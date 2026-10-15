"""Writers for task260's data files: the three ReadPoint exports, the case sets and panels, the RIS mammography
records, the examination codes, the location directory, the mobile units' visit schedule, the workforce list, and the
two data distractors (screening KPIs, readers' QA test-set scores)."""
from __future__ import annotations

import csv
import datetime as dt
import json

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import common as K

HDR = Font(name="Calibri", bold=True)
FILL = PatternFill("solid", fgColor="DCE6F1")
TITLE = Font(name="Calibri", bold=True, size=12)


def _csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\r\n")
        w.writerow(header)
        w.writerows(rows)
    return len(rows)


def _sheet(ws, title, header, rows, widths, note=None):
    start = 1
    if title:
        ws.append([title])
        ws.cell(row=1, column=1).font = TITLE
        if note:
            ws.append([note])
            ws.cell(row=2, column=1).alignment = Alignment(wrap_text=False)
        ws.append([])
        start = ws.max_row + 1
    ws.append(header)
    for c in ws[start]:
        c.font, c.fill = HDR, FILL
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A%d" % (start + 1)
    return start


# ------------------------------------------------------------------ ReadPoint exports

def session_numbers(log):
    """Sessions numbered 1 to 8 per reader in the order they were read."""
    first = {}
    for e in log:
        key = (e["uid"], e["cond"], e["session"])
        first[key] = min(first.get(key, e["opened"]), e["opened"])
    num = {}
    for uid in sorted({k[0] for k in first}):
        ks = sorted((k for k in first if k[0] == uid), key=lambda k: first[k])
        for i, k in enumerate(ks, 1):
            num[k] = i
    return num


def _ts41(t):
    return t.strftime("%Y-%m-%d %H:%M:%S")


def _ts42(t):
    return t.strftime("%Y-%m-%dT%H:%M:%S.") + "%03dZ" % (t.microsecond // 1000)


def write_export(path, st):
    S, log = st["S"], st["log"]
    num = session_numbers(log)
    rows = []
    if S["platform"] == "4.1":
        for e in sorted(log, key=lambda e: (e["uid"], e["opened"])):
            rows.append([S["code"], e["uid"], e["case"], "AIDED" if e["cond"] == "aided" else "UNAIDED",
                         num[(e["uid"], e["cond"], e["session"])], e["position"], _ts41(e["opened"]),
                         _ts41(e["submitted"]), e["pom"], "Y" if e["recall"] else "N", e["rev"]])
        header = ["project", "user_id", "case_code", "arm", "session_no", "position", "opened_utc", "submitted_utc",
                  "pom", "recall", "revision"]
    else:
        ws = {x["uid"]: "RP-WS-%02d" % (3 + i) for i, x in enumerate(st["panel"])}
        for e in sorted(log, key=lambda e: (e["opened"], e["uid"])):
            rows.append([S["code"], e["uid"], e["case"], "AI-assisted" if e["cond"] == "aided" else "Unaided",
                         num[(e["uid"], e["cond"], e["session"])], e["position"], _ts42(e["opened"]),
                         _ts42(e["submitted"]), "%.2f" % (e["pom"] / 100.0),
                         "Recall to assessment" if e["recall"] else "Return to routine recall", e["rev"] - 1,
                         ws[e["uid"]]])
        header = ["project_code", "reader_ref", "case_ref", "condition", "session", "seq", "started_at",
                  "completed_at", "malignancy_probability", "recall_decision", "edit_count", "workstation"]
    return _csv(path, header, rows)


def write_cases(path, studies):
    wb = Workbook()
    wb.remove(wb.active)
    for st in studies:
        S = st["S"]
        ws = wb.create_sheet(S["code"])
        rows = []
        for k, c in enumerate(st["L"]["cases"]):
            rows.append([st["code"][k], c["archive"], "Malignant" if c["cancer"] else "Normal", c["route"], c["months"],
                         c["side"], c["finding"], c["size"], c["grade"], c["density"], c["age"], c["year"]])
        rows.sort(key=lambda r: r[0])
        _sheet(ws, "%s study set: %s reader study" % (S["code"], S["tool"]),
               ["Case", "Archive episode", "Outcome", "Route", "Months to diagnosis", "Side", "Finding",
                "Invasive size (mm)", "Grade", "Breast density", "Age band", "Screen year"], rows,
               (10, 15, 11, 22, 19, 6, 14, 18, 18, 14, 10, 11),
               note="75 malignant and 225 normal examinations from the 2019 to 2022 screening archive. Outcome is the "
                    "reference standard: histology, or a normal screen at the next round.")
    wb.save(path)


def write_panels(path, studies):
    wb = Workbook()
    ws = wb.active
    ws.title = "Panels"
    rows = []
    for st in studies:
        S = st["S"]
        days = {}
        for e in st["log"]:
            days.setdefault((e["uid"], e["cond"]), set()).add(e["opened"].date())
        for x in st["panel"]:
            rd = x["reader"]
            fmt = lambda c: "; ".join(d.strftime("%d %b").lstrip("0") for d in sorted(days[(x["uid"], c)]))
            rows.append([S["code"], S["tool"], x["uid"], rd["code"], rd["role"], rd["years"],
                         "Unaided" if x["first"] == "unaided" else "With the tool", fmt("unaided"), fmt("aided")])
    _sheet(ws, "Reader study panels 2026", ["Study", "Tool", "ReadPoint user", "Reader", "Role",
                                            "Years reading screening mammograms", "Read first", "Unaided sessions",
                                            "Sessions with the tool"], rows,
           (12, 16, 15, 8, 36, 20, 14, 34, 34),
           note="Ten film readers per study, drawn within each role in proportion to the service's film readers;"
                " sessions with and without the tool at least four weeks apart.")
    wb.save(path)


# ------------------------------------------------------------------ the RIS

def write_ris(path, R):
    rows = []
    for s in sorted(R, key=lambda s: s["accession"]):
        t = s["utc"] if s["system"] == "native" else s["local"]
        rows.append([s["accession"], s["woman"], t.strftime("%Y-%m-%dT%H:%M:00Z"), s["loc"], s["exam"], s["status"],
                     "" if (s["system"] != "native" or s["images"] is None) else s["images"],
                     "" if (s["system"] != "native" or s["views"] is None) else s["views"], s["rg"]])
    return _csv(path, ["accession", "patient_ref", "study_datetime", "location_code", "exam_code", "status", "images",
                       "views", "radiographer"], rows)


def write_codes(path):
    return _csv(path, ["exam_code", "description", "programme"], [list(c[:3]) for c in K.EXAM_CODES])


def write_locations(path):
    locs = []
    for s in K.SITES:
        locs.append(dict(code=s["ris"], name="%s, screening" % s["name"], type="static screening", site=s["code"]))
        if s["assess"]:
            locs.append(dict(code=s["assess"], name="%s, assessment clinic" % s["name"], type="assessment clinic",
                             site=s["code"]))
        if s["sympt"]:
            locs.append(dict(code=s["sympt"], name="%s, symptomatic clinic" % s["name"], type="symptomatic clinic",
                             site=s["code"]))
    for code, host in sorted(K.MOBILE.items()):
        locs.append(dict(code=code, name="Mobile screening unit %s" % code[-1], type="mobile screening",
                         host_site=host))
    doc = {"system": "Merrowdale RIS", "table": "locations", "exported": "2026-10-05T08:47:12Z",
           "fields": {"site": "the screening site a location belongs to",
                      "host_site": "the screening site hosting a mobile unit at the time of export, whose film "
                                   "readers read its studies"},
           "sites": [dict(code=s["code"], name=s["name"], town=s["town"]) for s in K.SITES], "locations": locs}
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(doc, fh, indent=2)
        fh.write("\n")


def write_visits(path, V):
    wb = Workbook()
    ws = wb.active
    ws.title = "Programme 2025-26"
    names = {s["code"]: s["name"] for s in K.SITES}
    rows = [[v["unit"], v["vehicle"], names[v["host"]], v["practice"], v["start"], v["end"], v["days"]]
            for v in sorted(V, key=lambda v: (v["unit"], v["start"]))]
    _sheet(ws, "Mobile screening programme, October 2025 to September 2026, as delivered",
           ["Unit", "Vehicle", "Host site", "Location", "From", "To", "Screening days"], rows,
           (9, 10, 34, 30, 12, 12, 15))
    for r in range(4, ws.max_row + 1):
        for c in (5, 6):
            ws.cell(row=r, column=c).number_format = "dd/mm/yyyy"
    wb.save(path)


def write_workforce(path, rows_):
    wb = Workbook()
    ws = wb.active
    ws.title = "Breast screening staff"
    rows = [[w["emp"], w["name"], w["role"], w["reg"] or None, w["reads"], w["contract"], w["fte"],
             next(s["name"] for s in K.SITES if s["code"] == w["site"]), w["start"], w["end"]] for w in rows_]
    _sheet(ws, "Breast screening workforce: HR extract, 9 October 2026",
           ["Employee number", "Name", "Staff group", "Professional registration", "Reads films", "Contract", "FTE",
            "Base", "Start date", "End date"], rows, (17, 28, 34, 25, 12, 12, 6, 36, 12, 12),
           note="End dates include leaving dates notified to HR. Appointed staff appear from their offer acceptance.")
    for r in range(4, ws.max_row + 1):
        for c in (9, 10):
            ws.cell(row=r, column=c).number_format = "dd/mm/yyyy"
    wb.save(path)


# ------------------------------------------------------------------ distractors

def write_kpis(path):
    r = K.rng("kpis")
    wb = Workbook()
    ws = wb.active
    ws.title = "KPIs"
    rows = []
    for yr in ("2023-24", "2024-25", "2025-26"):
        for s in K.SITES:
            rows.append([yr, s["name"], round(float(r.uniform(67.5, 74.5)), 1),
                         round(float(r.uniform(3.6, 4.8)), 2), round(float(r.uniform(7.2, 9.8)), 1),
                         round(float(r.uniform(42.0, 58.0)), 1)])
    _sheet(ws, "Screening programme KPIs by screening year (April to March)",
           ["Screening year", "Site", "Uptake (%)", "Recall to assessment (%)",
            "Cancers detected per 1,000 screened", "Invasive cancers under 15 mm (%)"], rows,
           (15, 36, 12, 24, 33, 32),
           note="Women are counted at the site of their GP practice's screening round. Figures for 2025-26 are "
                "provisional.")
    wb.save(path)


def write_qa(path, RD):
    r = K.rng("qa")
    rows = []
    for x in sorted(RD, key=lambda x: x["code"]):
        cr = int(np.clip(round(r.normal(16.5 + 3 * x["skill"], 1.6)), 10, 20))
        nr = int(np.clip(round(r.normal(4.2 - 2 * x["skill"], 1.5)), 0, 12))
        rows.append([x["code"], "Self-assessment round 2025", "2025-%02d-%02d" % (int(r.integers(5, 12)), int(r.integers(1, 29))),
                     60, 20, cr, nr, "%.1f" % (100 * cr / 20), "%.1f" % (100 * (40 - nr) / 40)])
    return _csv(path, ["reader", "round", "read_on", "cases", "cancers", "cancers_recalled", "normals_recalled",
                       "sensitivity_pct", "specificity_pct"], rows)
