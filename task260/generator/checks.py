"""task260's generator checks: the input gates, the bundle against the generator's clean truth, the three rules
executed (R1 to R3c), every semantic fork row (SF1 to SF18), the ladder (every interval route and its pick, the
margins), the asks' devices, stops and cracker sheet, the separation line, the goldens against figures.json, and the
realism sweep."""
from __future__ import annotations

import collections
import datetime as dt
import json
import os
import re
import sys

import numpy as np
import pandas as pd

import common as K
import deliver as G
import load as L
import mrmc as M

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))


class Checks:
    def __init__(self):
        self.n = 0

    def __call__(self, label, cond, detail=""):
        if not cond:
            raise AssertionError("%s: %s" % (label, detail))
        self.n += 1


def pdf_text(path):
    from pypdf import PdfReader
    return "\n".join(p.extract_text() or "" for p in PdfReader(path).pages)


def trapezoid_auc(truth, score):
    """The area of the empirical ROC curve by the trapezoid rule over every distinct threshold."""
    th = np.unique(score)[::-1]
    tpr = [0.0] + [float((score[truth == 1] >= x).mean()) for x in th]
    fpr = [0.0] + [float((score[truth == 0] >= x).mean()) for x in th]
    return float(np.trapezoid(tpr, fpr))


def clean_mats(st):
    """The generator's own final readings per study, in the order deliver.study_matrix uses (readers by id, cases by
    code), with the reference standard."""
    out = {}
    for x in st["S"]["studies"]:
        Lt = x["L"]
        uids = [p["uid"] for p in x["panel"]]
        r_order = np.argsort(uids)
        codes = [x["code"][k] for k in range(len(Lt["cases"]))]
        c_order = np.argsort(codes)
        R = Lt["POM"][:, r_order][:, :, c_order].astype(float)
        out[x["S"]["sid"]] = (R, Lt["truth"][c_order], sorted(uids), sorted(codes))
    return out


def run(target, golden, st, meta_path):
    ck = Checks()
    out = st["res"]
    fig = out["figures"]
    T = out["T"]
    W = st["S"]
    tools = [S["tool"] for S in K.STUDIES]
    Q, V, A = "Ondrix Mammo", "Lumidra Detect", "Halvic Breast"

    # ---------------------------------------------------------------- input gates
    files = sorted(os.listdir(target))
    ck("gate: 10 or more files", len(files) >= 10, len(files))
    ck("gate: 3 or more formats", len({os.path.splitext(f)[1] for f in files}) >= 3, files)
    ck("gate: a file of 25,000 rows or more", len(T["ris"]) >= 25000, len(T["ris"]))
    meta = json.load(open(meta_path))
    ck("gate: two or more distractors named in metadata.json", len(meta["distractor_files"]) >= 2, meta)
    ck("gate: every distractor ships", all(f in files for f in meta["distractor_files"]), meta["distractor_files"])
    ck("gate: no file name says distractor, answer or issue", not any(re.search("distract|answer|issue|known", f, re.I)
                                                                         for f in files), files)
    ck("gate: every file non-empty", all(os.path.getsize(os.path.join(target, f)) > 200 for f in files))

    # ---------------------------------------------------------------- the bundle against the clean truth
    cm = clean_mats(st)
    for sid, (R, t, readers, cases) in out["mats"].items():
        Rc, tc, rc, cc = cm[sid]
        ck("bundle: %s readings of record equal the generator's final readings" % sid, np.array_equal(R, Rc))
        ck("bundle: %s reference standard equals the generator's" % sid, np.array_equal(t, tc))
        ck("bundle: %s ten readers and 300 cases (75 cancers)" % sid, R.shape == (2, 10, 300) and t.sum() == 75,
           (R.shape, t.sum()))
    for S in K.STUDIES:
        ex = T["exports"][S["sid"]]
        cs = set(T["cases"][S["sid"]]["Case"].astype(str))
        prac = ex[~ex["case"].isin(cs)]
        ck("bundle: %s practice cases open every session (five each) and are outside the study set" % S["sid"],
           len(prac) > 0 and prac.groupby(["reader", "session"]).size().eq(5).all()
           and prac["case"].str.contains("PRAC").all())
        sess = ex.groupby("reader")["session"].nunique()
        ck("bundle: %s eight sessions per reader" % S["sid"], (sess == 8).all(), sess.to_dict())
        fin = ex[ex["case"].isin(cs)].groupby(["reader", "case", "cond"]).size()
        ck("bundle: %s each reader read each case once per condition, resubmissions aside" % S["sid"],
           len(fin) == 6000 and fin.max() <= 2)
        ck("bundle: %s submissions follow their opening" % S["sid"], (ex["submitted"] > ex["opened"]).all())
        ck("bundle: %s every session before the export date" % S["sid"],
           ex["submitted"].max().date() < S["export"], ex["submitted"].max())
        wash = []
        for r_, g in ex.groupby("reader"):
            d = g.groupby("cond")["opened"].agg(["min", "max"])
            first = d["max"].min()
            second = d["min"].max()
            wash.append((second - first).total_seconds() / 86400)
        ck("bundle: %s at least four weeks between a reader's two conditions, by the clock" % S["sid"],
           min(wash) >= 28, min(wash))

    # ---------------------------------------------------------------- the call
    res = out["res"]
    ck("call: Halvic Breast", fig["call"] == A, fig["call"])
    ck("call: only Halvic's interval lies above zero", fig["qualifies"] == {V: False, Q: False, A: True}, fig["qualifies"])
    for k in tools:
        x1000 = res[k]["gain"] * 1000
        ck("rounding: %s's gain at least 0.15 thousandths from a rounding edge" % k, abs(x1000 - round(x1000)) <= 0.35,
           x1000)
        ck("boundary: %s's lower bound at least 0.005 from zero" % k, abs(res[k]["lo"]) >= 0.005, res[k]["lo"])
    ck("margin: Ondrix's gain at least 1.15 times Halvic's", res[Q]["gain"] >= 1.15 * res[A]["gain"])
    ck("margin: Lumidra's gain at least 1.15 times Halvic's", res[V]["gain"] >= 1.15 * res[A]["gain"])
    zs = {k: res[k]["ci"]["z"] for k in tools}
    crit = {k: res[k]["ci"]["crit"] for k in tools}
    ck("margin: Halvic at least 3.0 standard errors above zero, 0.9 above its critical value",
       zs[A] >= 3.0 and zs[A] >= crit[A] + 0.9, (zs[A], crit[A]))
    ck("margin: Ondrix at most 1.65 standard errors, 0.4 under its critical value", zs[Q] <= 1.65 and
       zs[Q] <= crit[Q] - 0.4, (zs[Q], crit[Q]))
    ck("margin: Lumidra at most 1.55 standard errors, 0.4 under its critical value", zs[V] <= 1.55 and
       zs[V] <= crit[V] - 0.4, (zs[V], crit[V]))

    # ---------------------------------------------------------------- the ladder: every route and its pick
    def pick(method, mats=None, keep="last"):
        r_, p_ = G.evaluate(T, method, keep=keep, mats=mats if mats is not None else out["mats"])
        return p_, r_
    routes = {}
    for m in ("fixed_dl", "fixed_jk", "boot_cases", "pooled", "paired_t", "signed_rank", "boot_readers", "or_jk",
              "or_dl", "dbm", "boot_both", "or_normal", "or_fpc", "or_one_sided", "or_bonferroni"):
        routes[m] = pick(m)
    for m in ("fixed_dl", "fixed_jk", "boot_cases"):
        ck("ladder R0: %s (the panel fixed) admits Ondrix and Halvic and picks Ondrix" % m,
           routes[m][0] == Q and routes[m][1][V]["lo"] < 0 and routes[m][1][A]["lo"] > 0, routes[m][0])
    ck("ladder R1: the pooled ROC curve admits all three and picks Lumidra",
       routes["pooled"][0] == V and all(routes["pooled"][1][k]["lo"] > 0 for k in tools), routes["pooled"][0])
    for m in ("paired_t", "signed_rank", "boot_readers"):
        ck("ladder R2: %s (the examinations fixed) picks Lumidra" % m, routes[m][0] == V, routes[m][0])
    ck("ladder R2: a paired t-test leaves Ondrix out", routes["paired_t"][1][Q]["lo"] < 0)
    for m in ("or_jk", "or_dl", "dbm", "boot_both"):
        ck("ladder R3 C1: %s (readers and examinations random) admits Halvic alone" % m,
           routes[m][0] == A and [k for k in tools if routes[m][1][k]["lo"] > 0] == [A], routes[m][0])
    ck("C1: DBM equals Obuchowski-Rockette with jackknife covariances", all(
        abs(routes["dbm"][1][k]["lo"] - routes["or_jk"][1][k]["lo"]) < 1e-9 for k in tools))
    ck("ladder: every wrong route names a tool other than Halvic", all(
        routes[m][0] != A for m in ("fixed_dl", "fixed_jk", "boot_cases", "pooled", "paired_t", "signed_rank",
                                    "boot_readers")))
    worse = fig["readers_worse"]
    ck("reason: Ondrix's gain swings from reader to reader (some readers worse, the spread wide)",
       worse[Q] >= 3 and np.std(res[Q]["A"][1] - res[Q]["A"][0], ddof=1) >= 0.06, worse[Q])
    ck("reason: every Lumidra reader did better with it", worse[V] == 0, worse[V])
    ck("reason: Lumidra's examination term dominates its interval (cov2 - cov3 over MS(T*R)/J)",
       10 * (res[V]["ci"]["cov2"] - res[V]["ci"]["cov3"]) > 3 * res[V]["ci"]["ms_tr"])

    # ---------------------------------------------------------------- the three rules on the clean truth (R1, R3)
    ck("R1: the readings ship one row per reader, case, condition and submission, over 6,000 per study",
       all(len(T["exports"][S["sid"]]) > 6000 for S in K.STUDIES))
    banned = re.compile(r"auc|area|gain|interval|signific|qualif|verdict|p_value|pvalue|lower|upper", re.I)
    heads = []
    for f in files:
        pth = os.path.join(target, f)
        if f.endswith(".csv"):
            heads += list(pd.read_csv(pth, nrows=0).columns)
        elif f.endswith(".xlsx"):
            from openpyxl import load_workbook
            wb = load_workbook(pth, read_only=True)
            for ws in wb.worksheets:
                for row in ws.iter_rows(min_row=1, max_row=5, values_only=True):
                    heads += [str(v) for v in row if v is not None and len(str(v)) < 40]
    ck("R1: no shipped column holds an area, gain, interval or verdict", not any(banned.search(h) for h in heads),
       [h for h in heads if banned.search(h)])
    r0, _ = pick("fixed_dl", mats=cm)
    rb, _ = pick("boot_cases", mats=cm)
    ck("R3a: the easy route on the clean truth (the examinations resampled, the panel fixed) files Ondrix",
       r0 == Q and rb == Q, (r0, rb))
    r3, _ = pick("or_jk", mats=cm)
    ck("R3b: the decisive step alone (readers and examinations random) on the clean truth files Halvic", r3 == A, r3)
    r2, _ = pick("paired_t", mats=cm)
    ck("R3b: the other one-factor correction (readers random alone) on the clean truth files Lumidra", r2 == V, r2)
    for S in K.STUDIES:
        ex = T["exports"][S["sid"]]
        ck("R3c: %s no reading row repeated" % S["sid"], not ex.duplicated(["reader", "case", "cond", "revision",
                                                                            "session"]).any())
        R_ = out["mats"][S["sid"]][0]
        sig = {tuple(R_[i, j]) for i in range(2) for j in range(10)}
        ck("R3c: %s no reader's ratings copied from another's" % S["sid"], len(sig) == 20)
    texts = {f: pdf_text(os.path.join(target, f)) for f in files if f.endswith(".pdf")}
    alltext = "\n".join(texts.values())
    ck("R3c and Gate G: no shipped document states a gain, an interval or a verdict for a tool",
       not re.search(r"(gain|interval)[^.]{0,40}\d\.\d{2,3}", alltext, re.I))
    ck("R3 not rule application: no filed sentence names a variance method or a random factor",
       not re.search(r"obuchowski|rockette|dbm|delong|bootstrap|jackknife|random effect|variance|t-test|readers? "
                     r"(as|are) random", alltext, re.I))

    # ---------------------------------------------------------------- semantic forks
    ck("SF1: the pooled ROC area gives another pick (Lumidra)", routes["pooled"][0] == V)
    for S in K.STUDIES:
        R_, t_, _, _ = out["mats"][S["sid"]]
        D_ = M.kernels(R_, t_)
        Am = M.aucs(D_)
        trap = np.array([[trapezoid_auc(t_, R_[i, j]) for j in range(10)] for i in range(2)])
        ck("SF2: %s the trapezoidal area of the empirical curve equals Mann-Whitney" % S["sid"],
           np.allclose(trap, Am, atol=1e-12))
    rec = {}
    for S in K.STUDIES:
        ex = T["exports"][S["sid"]]
        cs = T["cases"][S["sid"]]
        tr = dict(zip(cs["Case"].astype(str), (cs["Outcome"] == "Malignant").astype(int)))
        e = ex[ex["case"].isin(tr)].sort_values("revision").groupby(["reader", "case", "cond"]).tail(1)
        R2 = np.stack([e[e["cond"] == c].pivot(index="reader", columns="case", values="recall").astype(float).values
                       for c in ("unaided", "aided")])
        t2 = np.array([tr[c] for c in sorted(tr)])
        rec[S["tool"]] = M.aucs(M.kernels(R2, t2))
    ck("SF3: the recall decision as a two-point rating gives other areas (another quantity)", all(
        not np.allclose(rec[k], res[k]["A"]) for k in tools))
    pf, rf = pick("or_jk", keep="first", mats={S["sid"]: G.study_matrix(T, S["sid"], keep="first") for S in K.STUDIES})
    ck("SF4 R2: each case's first submission in place of the last, the same call", pf == A, pf)
    ck("SF4 R2: and the resubmissions move the gains (texture, not a device)", any(
        abs(rf[k]["gain"] - res[k]["gain"]) > 1e-6 for k in tools))
    ck("SF5 R2: no practice case in any study set", not any(
        T["cases"][S["sid"]]["Case"].astype(str).str.contains("PRAC").any() for S in K.STUDIES))
    raw3 = pd.read_csv(os.path.join(target, K.EXPORTS["S3"]), dtype=str)
    ck("SF6 R2: the 4.2 export's fractions are hundredths of the 0 to 100 scale",
       raw3["malignancy_probability"].str.fullmatch(r"[01]\.\d\d").all())
    for S in K.STUDIES:
        cs = T["cases"][S["sid"]]
        ck("SF7: %s interval cancers are malignant in the reference standard" % S["sid"],
           (cs.loc[cs["Route"] == "Interval cancer", "Outcome"] == "Malignant").all()
           and (cs["Route"] == "Interval cancer").sum() > 10)
    ck("SF8: one-sided 95 per cent, the same call", routes["or_one_sided"][0] == A)
    ck("SF9: Bonferroni at 98.3 per cent, the same call", routes["or_bonferroni"][0] == A)
    proto = texts[K.PROTOCOL]
    ck("SF10: the protocol files the populations", "for the service's film readers reading the service's screening "
       "examinations" in re.sub(r"\s+", " ", proto), proto[:200])
    ck("SF10: holding either population fixed gives Ondrix or Lumidra", routes["fixed_dl"][0] == Q and
       routes["paired_t"][0] == V)
    ck("SF11: every construction with both random gives the same call (OR jackknife and DeLong, DBM, two-way "
       "bootstrap, normal approximation)", all(routes[m][0] == A for m in ("or_jk", "or_dl", "dbm", "boot_both",
                                                                         "or_normal")))
    ck("SF12: a finite-population correction for 10 of 26 readers, the same call", routes["or_fpc"][0] == A)
    pan = T["panels"]
    for S in K.STUDIES:
        k = S["tool"]
        first = dict(zip(pan.loc[pan["Study"] == S["code"], "ReadPoint user"],
                         pan.loc[pan["Study"] == S["code"], "Read first"]))
        d = res[k]["A"][1] - res[k]["A"][0]
        grp = np.array([first[u] == "Unaided" for u in res[k]["readers"]])
        ck("SF13: %s half the panel read unaided first" % S["sid"], grp.sum() == 5)
        ck("SF13: %s the gain averaged within the two reading orders equals the gain" % S["sid"],
           abs((d[grp].mean() + d[~grp].mean()) / 2 - d.mean()) < 1e-12)
    group = lambda g: "radiologist" if "radiologist" in g else ("clinician" if "clinician" in g else "radiographer")
    wfr = T["workforce"][T["workforce"]["Reads films"] == "Y"]
    wreg = wfr["Professional registration"].astype(str).str.split().str[-1]
    now = wfr.assign(n=wreg)[pd.to_datetime(wfr["End date"]).isna() | (pd.to_datetime(wfr["End date"]) >= pd.Timestamp(
        "2026-10-09"))]
    now = now[pd.to_datetime(now["Start date"]) <= pd.Timestamp("2026-10-09")].drop_duplicates("n")
    share = now["Staff group"].map(group).value_counts(normalize=True).to_dict()
    for S in K.STUDIES:
        k = S["tool"]
        roles = dict(zip(pan.loc[pan["Study"] == S["code"], "ReadPoint user"],
                         pan.loc[pan["Study"] == S["code"], "Role"].map(group)))
        rr = [roles[u] for u in res[k]["readers"]]
        pshare = pd.Series(rr).value_counts(normalize=True).to_dict()
        w = np.array([share[x] / pshare[x] for x in rr])
        d = res[k]["A"][1] - res[k]["A"][0]
        wg = float((w * d).sum() / w.sum())
        bw = M.boot(res[k]["D"] if "D" in res[k] else M.kernels(*out["mats"][S["sid"]][:2]), 2000, 2027, wts=w)
        ck("SF17: %s weighted to the workforce's roles moves the gain by under 0.003 and keeps the verdict" % k,
           abs(wg - res[k]["gain"]) < 0.003 and (bw["lo"] > 0) == fig["qualifies"][k], (wg, res[k]["gain"], bw))
    ck("SF14: the largest gain across all tools is Lumidra's", max(tools, key=lambda k: res[k]["gain"]) == V)
    ck("SF15: the protocol keeps cost out of the clinical decision", "do not change it" in re.sub(r"\s+", " ", proto))
    ck("SF16: the hold is licensed and not reached", "keeps unaided double reading" in re.sub(r"\s+", " ", proto))
    for S in K.STUDIES:
        R_, t_, _, _ = out["mats"][S["sid"]]
        D_ = M.kernels(R_, t_)
        wc, wn = 0.0083 / 0.25, (1 - 0.0083) / 0.75          # each class reweighted to the screening cancer rate
        P_, N_ = wc * np.ones(D_.shape[2]), wn * np.ones(D_.shape[3])
        Aw = np.einsum("ijmn,m,n->ij", D_, P_, N_) / (P_.sum() * N_.sum())
        ck("SF18: %s reweighting the enriched set to the screening cancer rate leaves every area unchanged" % S["sid"],
           np.allclose(Aw, M.aucs(D_), atol=1e-12))

    # ---------------------------------------------------------------- the asks
    Rr = W["ris"]
    truth_site = collections.Counter()
    for s_ in Rr:
        if s_["copy"] or s_["status"] != "COMPLETED" or s_["exam"] not in ("MAMS", "MAMT"):
            continue
        if K.WINDOW[0] <= s_["local"].date() <= K.WINDOW[1]:
            truth_site[s_["host"] if s_["loc"].startswith("MBSU") else s_["loc"][:3]] += 1
    gold = fig["studies_per_site"]
    ck("Ask A: each site's count equals the generator's truth", gold == dict(truth_site), (gold, dict(truth_site)))
    ck("Ask A: no two sites share a count", len(set(gold.values())) == 4)
    ncopy = sum(1 for s_ in Rr if s_["copy"] and s_["status"] == "COMPLETED" and s_["exam"] in ("MAMS", "MAMT"))
    sg = out["s"]
    raw_n = len(G.screening_studies(T, dedupe="none"))
    ck("Ask A device D1: the golden removes exactly the history load's copies", raw_n - len(sg) == ncopy,
       (raw_n - len(sg), ncopy))
    cop_sites = collections.Counter(s_["site"] for s_ in Rr if s_["copy"])
    ck("Ask A device D1: every site has copies", all(cop_sites[c] >= 40 for c in K.SITE_ORDER), cop_sites)
    post = [s_ for s_ in Rr if s_["copy"] and s_["local"] >= dt.datetime(2026, 3, 29, 1)]
    ck("Ask A device D1: copies after 29 March sit an hour after their originals at every site",
       all(any(p_["site"] == c for p_ in post) for c in K.SITE_ORDER))
    imp = collections.Counter(s_["site"] for s_ in Rr if s_["kind"] == "implant" and not s_["copy"])
    ck("Ask A hazard: implant views registered as a second routine study at every site",
       all(imp[c] >= 60 for c in K.SITE_ORDER), imp)
    mu2_old = sum(1 for s_ in Rr if s_["loc"] == "MBSU2" and not s_["copy"] and s_["local"].date() <= K.MU2_OLD_LAST)
    ck("Ask A device D7: the replaced unit 2 vehicle's studies count to Corrowby Road", mu2_old > 1000, mu2_old)
    ck("Ask A device D7: the directory carries only the current host", next(
        x for x in T["locations"]["locations"] if x["code"] == "MBSU2")["host_site"] == "WDL")
    stops = {"no deduplication": G.per_site(G.screening_studies(T, dedupe="none")),
             "same woman, day, site and code": G.per_site(G.screening_studies(T, dedupe="same_day")),
             "exact time": G.per_site(G.screening_studies(T, dedupe="exact_time")),
             "the directory's host": G.per_site(G.screening_studies(T, dedupe="golden", hosts="directory"))}
    for nm, v in stops.items():
        off = [c for c in K.SITE_ORDER if v[c] != gold[c]]
        ck("Ask A stop (%s) misses its sites" % nm, len(off) >= 2, (nm, v))
    for dd in ("same_day", "exact_time", "none"):
        for hh in ("directory", "schedule"):
            v = G.per_site(G.screening_studies(T, dedupe=dd, hosts=hh))
            ck("Ask A cracker sheet: %s with %s hosts misses every site" % (dd, hh),
               all(v[c] != gold[c] for c in K.SITE_ORDER), v)
    att = set()
    for s_ in Rr:
        if not s_["copy"] and s_["status"] == "COMPLETED" and s_["kind"] == "screen" and \
                K.WINDOW[0] <= s_["local"].date() <= K.WINDOW[1]:
            att.add(s_["woman"])
    ck("Ask B: the golden episodes equal the generator's attendances", fig["episodes"] == len(att),
       (fig["episodes"], len(att)))
    late = sum(1 for s_ in Rr if s_["kind"] == "tr_later" and s_["exam"] == "MAMS" and not s_["copy"])
    rd_ = G.episodes(sg, "routine_days")
    ck("Ask B device D1: later-day technical recalls under the routine code", late >= 60, late)
    ck("Ask B device D1: counting routine study days counts them as episodes", rd_ == fig["episodes"] + late,
       (rd_, fig["episodes"], late))
    lq = re.sub(r"\s+", " ", texts[K.QUOTES["Lumidra"]])
    ck("Ask B pin: Lumidra's quote counts one episode per woman and screening round, with its routine study",
       "one woman's screening in a screening round" in lq and
       "in the period in which its routine screening study took place" in lq)
    ck("Ask B convergence: a 28-day window and every window from 29 to 1,000 days by calendar day give the golden",
       all(G.episodes(sg, "window", days=d) == fig["episodes"] for d in (28, 29, 45, 90, 182, 365, 1000)))
    wom = G.episodes(sg, "women")
    only_rep = sg.groupby("patient_ref")["exam_code"].agg(lambda x: set(x) == {"MAMT"})
    late_first = sg[sg["patient_ref"].isin(only_rep[only_rep].index)]["date"].max()
    ck("Ask B fork closed: every woman with a screening study counts the women whose only study is an early-October "
       "repeat (an episode begun before the window)", wom == fig["episodes"] + int(only_rep.sum()) and
       0 < only_rep.sum() < 30 and late_first < K.WINDOW[0] + dt.timedelta(days=21), (wom, int(only_rep.sum()),
                                                                                       late_first))
    clk = G.episodes(sg, "window", days=28, clock=True)
    ck("Ask B fork closed: a 28-day window by the clock splits a woman's round (the quote's round holds it)",
       clk != fig["episodes"] and abs(clk - fig["episodes"]) <= 3, clk)
    early = [s_ for s_ in Rr if s_["kind"] == "tr_later" and s_["exam"] == "MAMS" and not s_["copy"]
             and s_["local"].date() < K.WINDOW[0] + dt.timedelta(days=30)]
    ck("Ask B fairness: no routine-coded recall whose first study falls before the window", all(
        any(o["woman"] == e_["woman"] and o["kind"] == "screen" and o["local"].date() >= K.WINDOW[0] for o in Rr)
        for e_ in early))
    seats = {k: G.readers_in_post(T, k)[0] for k in ("registration", "registration_text", "employee", "name")}
    ck("Ask C: 27 film readers in post on 1 April 2027", seats["registration"] == 27 == fig["readers_in_post"], seats)
    ck("Ask C stops: employee numbers 28, names 28, registrations as written 28", seats["employee"] == 28 and
       seats["name"] == 28 and seats["registration_text"] == 28, seats)
    wf = T["workforce"]
    ck("Ask C hazards: a retirement on 31 March 2027 and a starter on 1 April 2027", (
        pd.to_datetime(wf["End date"]).dt.date == dt.date(2027, 3, 31)).sum() == 1 and (
        pd.to_datetime(wf["Start date"]).dt.date == dt.date(2027, 4, 1)).sum() == 1)
    c_g = G.costs(fig["studies_total"], fig["episodes"], fig["readers_in_post"])
    ck("Ask B: costs as figures.json", {k: G.pounds(v) for k, v in c_g.items()} == fig["cost"], fig["cost"])
    ck("Ask B rounding: no cost within 10p of a half pound", all(abs((v * 100) % 100 - 50) >= 10 for v in c_g.values()),
       c_g)
    cr = G.screening_studies(T, dedupe="same_day", hosts="directory")
    c_c = G.costs(sum(G.per_site(cr).values()), G.episodes(cr), seats["employee"])
    ck("Ask B cracker sheet: every tool's cost missed", all(G.pounds(c_c[k]) != fig["cost"][k] for k in tools),
       (c_c, fig["cost"]))

    # ---------------------------------------------------------------- the design note's figures recompute (H2)
    r2 = lambda x: round(float(x), 2)
    r3 = lambda x: round(float(x), 3)
    ck("H2 note: both-random z and critical values", [(r2(res[k]["ci"]["z"]), r2(res[k]["ci"]["crit"])) for k in
                                                       (A, Q, V)] == [(4.41, 2.10), (1.60, 2.09), (1.49, 1.96)])
    ck("H2 note: both-random intervals", [(r3(res[k]["lo"]), r3(res[k]["hi"])) for k in (V, Q, A)] ==
       [(-0.015, 0.109), (-0.013, 0.097), (0.014, 0.038)])
    fx = routes["fixed_dl"][1]
    ck("H2 note: Ondrix with the panel fixed, z 2.64, interval 0.011 to 0.073", r2(fx[Q]["ci"]["z"]) == 2.64 and
       (r3(fx[Q]["lo"]), r3(fx[Q]["hi"])) == (0.011, 0.073))
    pt_ = routes["paired_t"][1]
    ck("H2 note: across readers Lumidra z 5.32 (0.027 to 0.067), Ondrix z 1.93 (-0.007 to 0.091)",
       r2(pt_[V]["ci"]["z"]) == 5.32 and r2(pt_[Q]["ci"]["z"]) == 1.93 and (r3(pt_[V]["lo"]), r3(pt_[V]["hi"])) ==
       (0.027, 0.067) and (r3(pt_[Q]["lo"]), r3(pt_[Q]["hi"])) == (-0.007, 0.091))
    pl = routes["pooled"][1]
    ck("H2 note: pooled gains 0.061, 0.046, 0.026", [r3(pl[k]["gain"]) for k in (V, Q, A)] == [0.061, 0.046, 0.026])
    b2 = routes["boot_both"][1][A]
    ck("H2 note: Halvic's two-way bootstrap interval 0.007 to 0.045", (r3(b2["lo"]), r3(b2["hi"])) == (0.007, 0.045))
    dl = lambda v: [v[c] - gold[c] for c in K.SITE_ORDER]
    ck("H2 note: Ask A stops", dl(stops["no deduplication"]) == [149, 72, 227, 191] and
       dl(stops["same woman, day, site and code"]) == [-151, -90, -191, -194] and
       dl(stops["exact time"]) == [99, 72, 227, 191] and dl(stops["the directory's host"]) == [-5152, 5152, 0, 0],
       {k: dl(v) for k, v in stops.items()})
    crk = G.per_site(cr)
    ck("H2 note: the cracker's site counts", [crk[c] for c in K.SITE_ORDER] == [6326, 12300, 13062, 13564], crk)
    ck("H2 note: the cracker's costs", [G.pounds(c_c[k]) for k in (Q, V, A)] == [88108, 111085, 104200], c_c)
    st_b = {"routine_days": G.pounds(G.costs(fig["studies_total"], G.episodes(sg, "routine_days"),
                                             fig["readers_in_post"])[V]),
            "women": G.pounds(G.costs(fig["studies_total"], G.episodes(sg, "women"), fig["readers_in_post"])[V])}
    ck("H2 note: Lumidra's stops (routine study days 111,003; every woman 110,774)",
       st_b == {"routine_days": 111003, "women": 110774}, st_b)
    ck("H2 note: routine study days 44,386", rd_ == 44386, rd_)

    # ---------------------------------------------------------------- separation and distractors
    T2 = dict(T)
    for k in ("ris", "codes", "locations", "visits", "workforce", "panels"):
        T2.pop(k)
    _, p2 = G.evaluate(T2, "or_jk")
    ck("separation: the call reads no ask file (computed with every ask file removed)", p2 == A, p2)
    ck("separation: no ask device row in the call's files (the exports carry no RIS, workforce or location field)",
       not any(c in ("accession", "patient_ref", "location_code") for S in K.STUDIES
               for c in pd.read_csv(os.path.join(target, K.EXPORTS[S["sid"]]), nrows=0).columns))
    ck("distractors: the golden reads none of them (the loader opens none)", not any(
        x in json.dumps(list(T.keys())) for x in ("kpi", "qa", "summar")))
    summ = re.sub(r"\s+", " ", texts[K.PRODUCT_SUMMARIES])
    ck("wrong-basis distractor: the vendors' standalone areas rank Lumidra first, not the golden",
       "AUC 0.94" in summ and "Lumidra" in summ, summ[:300])

    # ---------------------------------------------------------------- goldens against figures.json
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(golden, K.DELIVERABLES[1]), data_only=True)
    tl = {r[0]: r for r in wb["Tools"].iter_rows(min_row=2, values_only=True) if r[0]}
    for k in tools:
        ck("golden xlsx: %s gain and verdict" % k, tl[k][6] == round(res[k]["gain"], 3) and
           tl[k][9] == ("Yes" if fig["qualifies"][k] else "No"), tl[k])
    st_ = {r[0]: r[1] for r in wb["Sites"].iter_rows(min_row=2, values_only=True) if r[0]}
    for c in K.SITE_ORDER:
        ck("golden xlsx: %s studies" % c, st_[G.SITE_NAMES[c]] == gold[c])
    co = {r[0]: r[6] for r in wb["Costs"].iter_rows(min_row=2, values_only=True) if r[0]}
    ck("golden xlsx: costs", all(co[k] == fig["cost"][k] for k in tools), co)
    paper = re.sub(r"\s+", " ", pdf_text(os.path.join(golden, K.DELIVERABLES[0])))
    ck("golden paper: the call and every gain", "license Halvic Breast" in paper and all(
        fig["gain_3dp"][k] in paper for k in tools))
    ck("golden paper: every cost and the reader count", all("{:,}".format(fig["cost"][k]) in paper for k in tools) and
       "27 film readers" in paper)
    ck("golden chart written", os.path.getsize(os.path.join(golden, K.DELIVERABLES[2])) > 20000)

    # ---------------------------------------------------------------- realism sweep
    ck("realism: site counts not round", all(v % 10 for v in gold.values()), gold)
    ck("realism: totals not round", fig["studies_total"] % 100 and fig["episodes"] % 100)
    ris = T["ris"]
    dates = pd.to_datetime(ris["study_datetime"]).dt.tz_localize(None).dt.date
    ck("realism: every RIS study date inside the window", dates.min() >= K.WINDOW[0] and dates.max() <= K.WINDOW[1])
    ck("realism: accession numbers unique", ris["accession"].is_unique)
    ck("realism: native records carry images and views, history records none", (
        ris.loc[ris["accession"].str.startswith("LR"), "views"] == "").all() and (
        ris.loc[ris["accession"].str.startswith("MW") & (ris["exam_code"] == "MAMS") & (ris["status"] == "COMPLETED"),
                "views"] != "").all())
    hours = pd.to_datetime(ris["study_datetime"]).dt.hour
    ck("realism: studies in working hours", hours.between(7, 18).all(), hours.describe())
    vis = T["visits"]
    ck("realism: visit blocks do not overlap within a unit", all(
        (g.sort_values("From")["From"].iloc[1:].values > g.sort_values("From")["To"].iloc[:-1].values).all()
        for _, g in vis.groupby("Unit")))
    sys.path.insert(0, os.path.join(REPO, ".claude", "skills", "fingerprint"))
    import business
    texts_all = alltext + "\n" + paper
    found = {"%s %s" % fl for fl in business.people_in_text(texts_all)}
    ck("realism: the documents name only the drawn people", found <= set(K.PEOPLE.values()) and len(found) >= 3, found)
    dash = re.compile("[–—]")
    ck("realism: no em or en dash in any document or golden", not dash.search(texts_all))
    for f in files:
        if f.endswith((".csv", ".json")):
            ck("realism: no em or en dash in %s" % f, not dash.search(open(os.path.join(target, f),
                                                                           encoding="utf-8").read()))
    ck("realism: no placeholder names", not re.search(r"John Doe|Jane Doe|Acme|Company A|Lorem", texts_all + json.dumps(
        list(wf["Name"]))))
    fr = wf[wf["Reads films"] == "Y"]
    regn = fr["Professional registration"].astype(str).str.split().str[-1]
    ck("realism: no two different film readers share a name (a repeated name is one person's two posts)",
       (fr.assign(n=regn).groupby("Name")["n"].nunique() == 1).all())
    ck("Ask C hazard and device: two registration numbers on two rows each, the changed post and the bank contract",
       (fr.assign(n=regn).groupby("n").size() == 2).sum() == 2)
    bk = wf[wf["Contract"] == "Bank"]
    ck("Ask C no marker: every bank row has no contracted hours and its registration without the prefix",
       bk["FTE"].isna().all() and not bk["Professional registration"].astype(str).str.startswith(("GMC", "HCPC")).any()
       and len(bk) >= 5)
    nm_reg = wf.assign(n=wf["Professional registration"].astype(str).str.split().str[-1])
    twins = nm_reg[nm_reg["n"].str.len() > 4].groupby("n")["Contract"].agg(lambda x: sorted(x) == ["Bank", "Substantive"])
    ck("Ask C no marker: a double contract on a reader and on a non-reader", twins.sum() == 2, twins[twins].index.tolist())
    ck("realism: GMC numbers only for doctors, HCPC for radiographers", all(
        (str(g).startswith("GMC") and grp in ("Consultant radiologist", "Breast clinician")) or
        (str(g).startswith("HCPC") and "radiographer" in grp) or not str(g).startswith(("GMC", "HCPC"))
        for g, grp in zip(fr["Professional registration"], fr["Staff group"])))
    print("checks: %d assertions passed" % ck.n)
    return ck.n
