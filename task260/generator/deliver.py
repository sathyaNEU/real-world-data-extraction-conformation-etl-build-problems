"""The golden, computed from the pack alone: each tool's gain in the panel readers' average area under the empirical
ROC curve and its interval with readers and examinations both random (Obuchowski-Rockette, jackknife covariance,
Hillis's degrees of freedom), the verdicts and the pick; then the asks (screening studies per site, film readers in
post on 1 April 2027, each tool's first-year cost) and the three deliverables."""
from __future__ import annotations

import datetime as dt
from decimal import ROUND_HALF_UP, Decimal

import numpy as np
import pandas as pd

import common as K
import load as L
import mrmc as M

BOOT_REPS = 2000
BOOT_SEED = 2027
SITE_NAMES = {s["code"]: s["name"] for s in K.SITES}


def pounds(x):
    return int(Decimal(repr(round(x, 6))).quantize(Decimal(1), rounding=ROUND_HALF_UP))


# ------------------------------------------------------------------ the call

def study_matrix(T, sid, keep="last"):
    """(2, J, K) probabilities of malignancy (unaided, aided; the panel's readers; the study set's cases) and the
    reference standard. keep: 'last' (the reading of record) or 'first' (each case's first submission)."""
    ex = T["exports"][sid]
    cs = T["cases"][sid]
    truth_of = dict(zip(cs["Case"].astype(str), (cs["Outcome"] == "Malignant").astype(int)))
    ex = ex[ex["case"].isin(truth_of)].sort_values(["reader", "case", "cond", "revision", "opened"])
    ex = ex.groupby(["reader", "case", "cond"]).tail(1) if keep == "last" else \
        ex.groupby(["reader", "case", "cond"]).head(1)
    readers, cases = sorted(ex["reader"].unique()), sorted(truth_of)
    ri, ci = {r: i for i, r in enumerate(readers)}, {c: i for i, c in enumerate(cases)}
    R = np.full((2, len(readers), len(cases)), np.nan)
    for i, cond in enumerate(("unaided", "aided")):
        sub = ex[ex["cond"] == cond]
        R[i, sub["reader"].map(ri).values, sub["case"].map(ci).values] = sub["pom"].values
    assert not np.isnan(R).any(), sid
    return R, np.array([truth_of[c] for c in cases]), readers, cases


def interval(R, t, method):
    D = M.kernels(R, t)
    A = M.aucs(D)
    if method == "or_jk":
        ci = M.or_test(A, M.jackknife_cov(D))
    elif method == "or_dl":
        ci = M.or_test(A, M.delong_cov(D))
    elif method == "or_normal":
        ci = M.or_test(A, M.jackknife_cov(D), normal=True)
    elif method == "or_fpc":
        ci = M.or_test(A, M.jackknife_cov(D), reader_fpc=(R.shape[1], 26))
    elif method == "or_one_sided":
        ci = M.or_test(A, M.jackknife_cov(D), level=0.90)
    elif method == "or_bonferroni":
        ci = M.or_test(A, M.jackknife_cov(D), level=1 - 0.05 / 3)
    elif method == "dbm":
        ci = M.dbm(D)
    elif method == "fixed_dl":
        ci = M.fixed_readers(A, M.delong_cov(D))
    elif method == "fixed_jk":
        ci = M.fixed_readers(A, M.jackknife_cov(D))
    elif method == "paired_t":
        ci = M.paired_t(A)
    elif method == "signed_rank":
        p = M.signed_rank_p(A)
        ci = dict(lo=1.0 if (p < 0.05 and (A[1] - A[0]).mean() > 0) else -1.0, hi=1.0, p=p)
    elif method == "boot_both":
        ci = M.boot(D, BOOT_REPS, BOOT_SEED)
    elif method == "boot_cases":
        ci = M.boot(D, BOOT_REPS, BOOT_SEED, readers=False)
    elif method == "boot_readers":
        ci = M.boot(D, BOOT_REPS, BOOT_SEED, cases=False)
    elif method == "pooled":
        ci = M.pooled(R, t)
        return dict(gain=float(ci["diff"]), lo=float(ci["lo"]), hi=float(ci["hi"]), A=A, ci=ci)
    else:
        raise ValueError(method)
    return dict(gain=float(A[1].mean() - A[0].mean()), lo=float(ci["lo"]), hi=float(ci["hi"]), A=A, ci=ci)


def evaluate(T, method="or_jk", keep="last", mats=None):
    """Each tool's gain and interval under one interval method, the qualifying tools and the pick."""
    out = {}
    for S in K.STUDIES:
        R, t, readers, cases = mats[S["sid"]] if mats else study_matrix(T, S["sid"], keep)
        out[S["tool"]] = dict(interval(R, t, method), readers=readers, sid=S["sid"])
    qual = [k for k, v in out.items() if v["lo"] > 0]
    pick = max(qual, key=lambda k: out[k]["gain"]) if qual else None
    return out, pick


# ------------------------------------------------------------------ the asks

def screening_studies(T, dedupe="golden", hosts="schedule"):
    """Completed screening studies (routine and technical repeat) with study dates in the 12 months to 30 September
    2026, one row per study, with the site whose film readers read them.
    dedupe: 'golden' (a legacy record matching a native one for the same woman, location, code and day at the same
    clock or exactly an hour later is the migration's copy), 'same_day' (one study per woman, day, location and code),
    'exact_time' (identical woman, location, code and time), 'none'.
    hosts: 'schedule' (a mobile unit's host on the day, from the visit schedule) or 'directory' (its current host)."""
    ris = T["ris"]
    s = ris[(ris["status"] == "COMPLETED") & ris["exam_code"].isin(["MAMS", "MAMT"])].copy()
    s["t"] = pd.to_datetime(s["study_datetime"]).dt.tz_localize(None)
    s["date"] = s["t"].dt.date
    s = s[(s["date"] >= K.WINDOW[0]) & (s["date"] <= K.WINDOW[1])]
    if dedupe == "golden":
        leg = s[s["accession"].str.startswith("LR")]
        nat = s[s["accession"].str.startswith("MW")]
        m = leg.reset_index().merge(nat[["patient_ref", "location_code", "exam_code", "date", "t"]],
                                    on=["patient_ref", "location_code", "exam_code", "date"], suffixes=("", "_n"))
        lag = (m["t"] - m["t_n"]).dt.total_seconds() / 60
        copies = set(m.loc[lag.isin([0.0, 60.0]), "index"])
        s = s.drop(index=list(copies))
    elif dedupe == "same_day":
        s = s.drop_duplicates(["patient_ref", "date", "location_code", "exam_code"])
    elif dedupe == "exact_time":
        s = s.drop_duplicates(["patient_ref", "location_code", "exam_code", "t"])
    loc = {x["code"]: x for x in T["locations"]["locations"]}
    vis = T["visits"]
    name_to_code = {x["name"]: x["code"] for x in T["locations"]["sites"]}
    blocks = [(str(r["Unit"]), pd.Timestamp(r["From"]).date(), pd.Timestamp(r["To"]).date(),
               name_to_code[r["Host site"]]) for _, r in vis.iterrows()]

    def site_of(code, day):
        x = loc[code]
        if x["type"] != "mobile screening":
            return x["site"]
        if hosts == "directory":
            return x["host_site"]
        hit = [h for u, a, b, h in blocks if u == code and a <= day <= b]
        assert len(hit) == 1, (code, day, hit)
        return hit[0]
    s["site"] = [site_of(c, d) for c, d in zip(s["location_code"], s["date"])]
    return s


def per_site(s):
    return {c: int((s["site"] == c).sum()) for c in K.SITE_ORDER}


def episodes(s, how="golden", days=None, clock=False):
    """Screening episodes in the window (Lumidra's quote: one woman's screening in a screening round, technical
    repeats included, counted in the period of its routine study). golden: one per woman with a routine study in the
    window (the round is three years, so a woman has at most one in twelve months; a woman whose only study in the
    window is a technical repeat began her episode before it).
    'routine_days': one per woman and day with a routine study; 'women': every woman with a screening study;
    'window': a routine study opens an episode unless within `days` of the one that opened her current episode, by
    calendar day or, with clock, by the time stamp."""
    if how == "routine_days":
        r = s[s["exam_code"] == "MAMS"]
        return int(r.drop_duplicates(["patient_ref", "date"]).shape[0])
    if how == "women":
        return int(s["patient_ref"].nunique())
    if how == "golden":
        return int(s.loc[s["exam_code"] == "MAMS", "patient_ref"].nunique())
    n = 0
    for _, g in s.sort_values("t").groupby("patient_ref", sort=False):
        start = None
        for code, day, t in zip(g["exam_code"], g["date"], g["t"]):
            gap = None if start is None else ((t - start).total_seconds() / 86400 if clock else (day - start).days)
            if code == "MAMS" and (gap is None or gap > days):
                start = t if clock else day
                n += 1
    return n


def readers_in_post(T, key="registration", day=K.CONTRACT_START):
    w = T["workforce"]
    f = w[w["Reads films"] == "Y"]
    start = pd.to_datetime(f["Start date"]).dt.date
    end = pd.to_datetime(f["End date"]).dt.date
    on = (start <= day) & (end.isna() | (end >= day))
    if key == "registration":
        reg = f["Professional registration"].astype(str).str.replace(r"^(GMC|HCPC)\s+", "", regex=True)
        return int(reg[on].nunique()), f.loc[on].assign(reg_number=reg[on])
    col = {"registration_text": "Professional registration", "employee": "Employee number", "name": "Name"}[key]
    return int(f.loc[on, col].nunique()), f.loc[on]


def costs(n_studies, n_episodes, n_seats, n_sites=4):
    P = K.PRICES
    q = max(P["Ondrix"]["per_study"] * n_studies, P["Ondrix"]["minimum"]) + P["Ondrix"]["one_off"]
    v = P["Lumidra"]["per_episode"] * n_episodes + P["Lumidra"]["per_reader"] * n_seats + P["Lumidra"]["one_off"]
    a = P["Halvic"]["per_seat"] * n_seats + P["Halvic"]["per_site"] * n_sites + P["Halvic"]["one_off"]
    return {"Ondrix Mammo": q, "Lumidra Detect": v, "Halvic Breast": a}


# ------------------------------------------------------------------ everything

def compute(target):
    T = L.read(target)
    mats = {S["sid"]: study_matrix(T, S["sid"]) for S in K.STUDIES}
    res, pick = evaluate(T, "or_jk", mats=mats)
    s = screening_studies(T)
    sites = per_site(s)
    ep = episodes(s)
    seats, seat_rows = readers_in_post(T)
    cost = costs(sum(sites.values()), ep, seats)
    tools = [S["tool"] for S in K.STUDIES]
    fig = {"call": pick,
           "gain": {k: round(res[k]["gain"], 6) for k in tools},
           "gain_3dp": {k: "%.3f" % res[k]["gain"] for k in tools},
           "interval": {k: [round(res[k]["lo"], 4), round(res[k]["hi"], 4)] for k in tools},
           "qualifies": {k: bool(res[k]["lo"] > 0) for k in tools},
           "unaided": {k: round(float(res[k]["A"][0].mean()), 4) for k in tools},
           "aided": {k: round(float(res[k]["A"][1].mean()), 4) for k in tools},
           "readers_worse": {k: int((res[k]["A"][1] < res[k]["A"][0]).sum()) for k in tools},
           "readers_better": {k: int((res[k]["A"][1] > res[k]["A"][0]).sum()) for k in tools},
           "df": {k: round(float(res[k]["ci"]["df"]), 1) for k in tools},
           "panel_readers": 10, "examinations_per_study": 300, "cancers_per_study": 75, "interval_level_pct": 95,
           "studies_per_site": sites, "studies_total": int(sum(sites.values())), "episodes": int(ep),
           "readers_in_post": seats, "cost": {k: pounds(v) for k, v in cost.items()}}
    return dict(figures=fig, T=T, res=res, mats=mats, s=s, seat_rows=seat_rows, pick=pick)


# ------------------------------------------------------------------ deliverables

def write_xlsx(path, out):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    f, res, T = out["figures"], out["res"], out["T"]
    hdr, fill = Font(name="Calibri", bold=True), PatternFill("solid", fgColor="DCE6F1")
    wb = Workbook()
    ws = wb.active
    ws.title = "Tools"
    ws.append(["Tool", "Reader study", "Panel readers", "Examinations", "Readers' average area, unaided",
               "Readers' average area, with the tool", "Gain", "95% interval, lower", "95% interval, upper",
               "Meets the protocol", "Readers who did better with the tool", "Why not"])
    for c in ws[1]:
        c.font, c.fill = hdr, fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
    why = {"Ondrix Mammo": "The gain differs widely from reader to reader (%d of 10 read worse with it), so it is not "
                             "shown for the service's film readers." % f["readers_worse"]["Ondrix Mammo"],
           "Lumidra Detect": "The gain varies from examination to examination, the same way for every reader, so it is "
                             "not shown for the service's screening examinations.",
           "Halvic Breast": None}
    for S in K.STUDIES:
        k = S["tool"]
        ws.append([k, S["code"], 10, "300 (75 cancers, 225 normal)", f["unaided"][k], f["aided"][k],
                   round(res[k]["gain"], 3), round(res[k]["lo"], 3), round(res[k]["hi"], 3),
                   "Yes" if f["qualifies"][k] else "No", f["readers_better"][k], why[k]])
        for col in (5, 6):
            ws.cell(row=ws.max_row, column=col).number_format = "0.000"
        for col in (7, 8, 9):
            ws.cell(row=ws.max_row, column=col).number_format = "0.000"
        ws.cell(row=ws.max_row, column=12).alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFGHIJKL", (16, 12, 8, 24, 16, 16, 8, 11, 11, 10, 14, 60)):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "B2"
    wr = wb.create_sheet("Readers")
    wr.append(["Tool", "ReadPoint user", "Reader", "Role", "Area unaided", "Area with the tool", "Change"])
    for c in wr[1]:
        c.font, c.fill = hdr, fill
    pan = T["panels"]
    for S in K.STUDIES:
        k = S["tool"]
        A = res[k]["A"]
        for j, uid in enumerate(res[k]["readers"]):
            p = pan[(pan["Study"] == S["code"]) & (pan["ReadPoint user"] == uid)].iloc[0]
            wr.append([k, uid, p["Reader"], p["Role"], round(float(A[0, j]), 3), round(float(A[1, j]), 3),
                       round(float(A[1, j] - A[0, j]), 3)])
            for col in (5, 6, 7):
                wr.cell(row=wr.max_row, column=col).number_format = "0.000"
    for col, w in zip("ABCDEFG", (16, 15, 8, 36, 13, 18, 9)):
        wr.column_dimensions[col].width = w
    wr.freeze_panes = "A2"
    wsite = wb.create_sheet("Sites")
    wsite.append(["Site", "Screening studies sent for reading, 1 Oct 2025 to 30 Sep 2026",
                  "Of which from mobile units"])
    for c in wsite[1]:
        c.font, c.fill = hdr, fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
    s = out["s"]
    mob = s["location_code"].str.startswith("MBSU")
    for code in K.SITE_ORDER:
        wsite.append([SITE_NAMES[code], f["studies_per_site"][code], int(((s["site"] == code) & mob).sum())])
    wsite.append(["All sites", f["studies_total"], int(mob.sum())])
    wsite.cell(row=wsite.max_row, column=1).font = hdr
    for r in range(2, wsite.max_row + 1):
        for col in (2, 3):
            wsite.cell(row=r, column=col).number_format = "#,##0"
    for col, w in zip("ABC", (36, 30, 16)):
        wsite.column_dimensions[col].width = w
    wsite.row_dimensions[1].height = 32
    wc = wb.create_sheet("Costs")
    wc.append(["Tool", "Charged on", "Volume", "Unit charge (GBP)", "Annual charges (GBP)", "One-off (GBP)",
               "First contract year, excluding VAT (GBP)"])
    for c in wc[1]:
        c.font, c.fill = hdr, fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
    P = K.PRICES
    q_annual = max(P["Ondrix"]["per_study"] * f["studies_total"], P["Ondrix"]["minimum"])
    wc.append(["Ondrix Mammo", "screening studies received for analysis", f["studies_total"],
               P["Ondrix"]["per_study"], round(q_annual, 2), P["Ondrix"]["one_off"], f["cost"]["Ondrix Mammo"]])
    wc.append(["Lumidra Detect", "screening episodes; film reader accounts",
               "{:,} episodes; {:d} accounts".format(f["episodes"], f["readers_in_post"]),
               "%.2f per episode; %d per account" % (P["Lumidra"]["per_episode"], P["Lumidra"]["per_reader"]),
               round(P["Lumidra"]["per_episode"] * f["episodes"] + P["Lumidra"]["per_reader"] * f["readers_in_post"],
                     2), P["Lumidra"]["one_off"], f["cost"]["Lumidra Detect"]])
    wc.append(["Halvic Breast", "film reader seats; screening sites", "%d seats; 4 sites" % f["readers_in_post"],
               "{:,} per seat; {:,} per site".format(P["Halvic"]["per_seat"], P["Halvic"]["per_site"]),
               P["Halvic"]["per_seat"] * f["readers_in_post"] + 4 * P["Halvic"]["per_site"], P["Halvic"]["one_off"],
               f["cost"]["Halvic Breast"]])
    for r in range(2, wc.max_row + 1):
        for col in (5, 6, 7):
            wc.cell(row=r, column=col).number_format = "#,##0" if col == 7 else "#,##0.00"
        if isinstance(wc.cell(row=r, column=3).value, int):
            wc.cell(row=r, column=3).number_format = "#,##0"
    for col, w in zip("ABCDEFG", (16, 34, 16, 20, 18, 12, 22)):
        wc.column_dimensions[col].width = w
    wc.row_dimensions[1].height = 32
    ww = wb.create_sheet("Film readers")
    ww.append(["Staff group", "Film readers in post on 1 April 2027"])
    for c in ww[1]:
        c.font, c.fill = hdr, fill
    rows = out["seat_rows"].drop_duplicates("reg_number")
    for g, n in rows.groupby("Staff group").size().items():
        ww.append([g, int(n)])
    ww.append(["All film readers", f["readers_in_post"]])
    ww.cell(row=ww.max_row, column=1).font = hdr
    ww.column_dimensions["A"].width = 36
    ww.column_dimensions["B"].width = 34
    wn = wb.create_sheet("Notes")
    notes = [
        ("Measure", "Each reader's area under the empirical ROC curve of their probability-of-malignancy ratings "
                    "against the study set's reference standard (ties counted half), the reading of record being the "
                    "reader's last submission; the gain is the panel's average area with the tool minus its average "
                    "area unaided (protocol 3)."),
        ("Interval", "Two-sided 95 per cent, with the panel's readers and the study's examinations both treated as "
                     "samples of the service's film readers and screening examinations (protocol 4.1): "
                     "Obuchowski-Rockette analysis with jackknife covariances and Hillis's degrees of freedom, the "
                     "same as DBM; DeLong covariances and a bootstrap of readers and examinations give the same "
                     "verdicts."),
        ("Studies", "Completed routine and technical repeat screening studies dated 1 October 2025 to 30 September "
                    "2026. Records the legacy system's history load added for studies already registered in the new "
                    "system during each site's parallel running are counted once; implant views registered as a "
                    "second study are counted. A mobile unit's studies count to the site hosting it on the day "
                    "(mobile visit schedule): unit 2's earlier vehicle to Corrowby Road until 29 May 2026."),
        ("Episodes", "One per woman and screening round, counted with its routine study (Lumidra's quote): technical "
                     "repeats, including a later-day recall booked under the routine code, belong to it, and a "
                     "repeat in early October 2025 belongs to an episode that began before the 12 months."),
        ("Film readers", "Staff who read films, in post on 1 April 2027 (start on or before, end on or after), "
                         "counted once per professional registration number (the bank system records the number "
                         "without the GMC or HCPC prefix): one reader holds a substantive and a bank contract."),
        ("Prepared", "AI evaluation lead, %s." % K.DELIVERY_DATE.strftime("%d %B %Y").lstrip("0")),
    ]
    for a, b in notes:
        wn.append([a, b])
        wn.cell(row=wn.max_row, column=1).font = hdr
        wn.cell(row=wn.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    wn.column_dimensions["A"].width = 14
    wn.column_dimensions["B"].width = 110
    wb.save(path)


def write_chart(path, out):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    f, res = out["figures"], out["res"]
    tools = [S["tool"] for S in K.STUDIES][::-1]
    fig, ax = plt.subplots(figsize=(9.4, 4.9), dpi=110)
    for y, k in enumerate(tools):
        ok = f["qualifies"][k]
        col = "#1f6f3a" if ok else "#9a4b2f"
        ax.plot([res[k]["lo"], res[k]["hi"]], [y, y], color=col, linewidth=2.4, solid_capstyle="butt", zorder=2)
        ax.scatter([res[k]["gain"]], [y], s=70, color=col, zorder=3)
        d = res[k]["A"][1] - res[k]["A"][0]
        ax.scatter(d, np.full(len(d), y + 0.22), s=12, color="#7a8691", alpha=0.8, zorder=1,
                   label="One panel reader's change" if y == 0 else None)
        ax.annotate("%.3f  [%.3f to %.3f]" % (res[k]["gain"], res[k]["lo"], res[k]["hi"]), (res[k]["hi"], y),
                    textcoords="offset points", xytext=(8, -4), fontsize=8.6, color=col)
    ax.axvline(0, color="#1d2329", linewidth=1.0, zorder=0)
    ax.annotate("no gain", (0, len(tools) - 0.45), textcoords="offset points", xytext=(4, 0), fontsize=8,
                color="#1d2329")
    ax.set_yticks(range(len(tools)))
    ax.set_yticklabels(["%s\n%s" % (k, "meets the protocol" if f["qualifies"][k] else "does not meet it")
                        for k in tools], fontsize=9)
    ax.set_ylim(-0.6, len(tools) - 0.3)
    lo = min(min(res[k]["lo"] for k in tools), min((res[k]["A"][1] - res[k]["A"][0]).min() for k in tools))
    hi = max(max(res[k]["hi"] for k in tools), max((res[k]["A"][1] - res[k]["A"][0]).max() for k in tools))
    ax.set_xlim(lo - 0.01, hi + 0.07)
    ax.set_xlabel("Gain in the readers' average area under the ROC curve (with the tool minus unaided)")
    ax.set_title("Only Halvic Breast's gain holds for the service's readers and examinations\n"
                 "Gain and two-sided 95% interval, readers and examinations both random\n"
                 "Three 2026 reader studies, 10 readers and 300 examinations each", fontsize=10.2, loc="left")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.legend(frameon=False, fontsize=8, loc="lower right")
    fig.text(0.01, 0.01, "Source: ReadPoint exports of the MBS-AI-01 to 03 reader studies; Obuchowski-Rockette "
                         "analysis (jackknife covariance). %s, AI evaluation." % K.ORG, fontsize=7.2, color="#5b6670")
    fig.tight_layout(rect=(0, 0.04, 1, 1))
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)


def write_paper(path, out):
    import docs as D
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Spacer
    f, res = out["figures"], out["res"]
    lead = K.PEOPLE["lead"]
    q, v, a = "Ondrix Mammo", "Lumidra Detect", "Halvic Breast"
    g = lambda k: "%.3f" % res[k]["gain"]
    ci = lambda k: "%.3f to %.3f" % (res[k]["lo"], res[k]["hi"])
    story = [Paragraph("AI reading support for 2027: license Halvic Breast", D.H1),
             Paragraph("To the AI evaluation panel (chair: %s, clinical director). From %s, AI evaluation lead. %s. "
                       "For the panel's decision on %s." % (
                           K.PEOPLE["director"], lead, K.DELIVERY_DATE.strftime("%d %B %Y").lstrip("0"),
                           K.PANEL_DATE.strftime("%d %B %Y").lstrip("0")), D.META),
             Paragraph("Recommendation", D.H2),
             Paragraph("License Halvic Breast for two years from 1 April 2027. Of the three tools trialled, it is the "
                       "only one whose gain in the readers' average area under the ROC curve has a 95 per cent "
                       "interval wholly above zero for the service's film readers reading the service's screening "
                       "examinations (protocol 4.1): a gain of %s (interval %s). Ondrix Mammo (%s) and Lumidra "
                       "Detect (%s) show larger average gains in their studies, but neither is shown to hold for the "
                       "service's readers and examinations, so under protocol 4.2 Halvic Breast is the tool to "
                       "license." % (g(a), ci(a), g(q), g(v)), D.BODY),
             D._table([["Tool", "Gain", "95% interval", "Meets the protocol", "Readers better with the tool"]] + [
                 [k, g(k), ci(k), "Yes" if f["qualifies"][k] else "No", "%d of 10" % f["readers_better"][k]]
                 for k in (v, q, a)], (40, 18, 36, 30, 42)),
             Paragraph("Source: ReadPoint exports of reader studies MBS-AI-01 (Lumidra Detect), MBS-AI-02 (Ondrix "
                       "Mammo) and MBS-AI-03 (Halvic Breast); intervals by Obuchowski-Rockette analysis with jackknife "
                       "covariances, readers and examinations random.", D.NOTE),
             Spacer(1, 3 * mm),
             Paragraph("Why not Ondrix or Lumidra", D.H2),
             Paragraph("Each study measured one panel of ten readers on one set of 300 examinations, so a tool's gain "
                       "carries two kinds of sampling: which readers sat on the panel and which examinations were in "
                       "the set. The decision is about every film reader we have and will have, reading the "
                       "examinations we will screen, so the interval has to carry both.", D.BODY),
             Paragraph("<b>Ondrix Mammo.</b> Its gain swings from reader to reader: %d of the ten read worse with it "
                       "and others gained a lot. Holding the panel fixed and resampling only the examinations (a "
                       "bootstrap of the study set, or DeLong's test) puts its interval clear of zero and would pick "
                       "it as the largest qualifying gain; once the readers are treated as a sample of ours, its "
                       "interval runs from %s." % (f["readers_worse"][q], ci(q)), D.BODY),
             Paragraph("<b>Lumidra Detect.</b> Every reader on its panel did better with it, but its gain varies from "
                       "examination to examination, up on some and down on others, the same way for every reader. A "
                       "test across the ten readers alone (each reader as their own control) puts its interval clear "
                       "of zero and would pick it; once the examinations are treated as a sample of ours, its interval "
                       "runs from %s." % ci(v), D.BODY),
             Paragraph("<b>Halvic Breast.</b> A smaller, steady gain: its interval stays above zero with readers and "
                       "examinations both random, under the Obuchowski-Rockette analysis (jackknife or DeLong "
                       "covariances), DBM, and a bootstrap of readers and examinations together.", D.BODY),
             Paragraph("Volumes, film readers and first-year costs", D.H2),
             Paragraph("The service's PACS sent %s screening studies for reading in the 12 months to 30 September 2026 "
                       "(%s), %s screening episodes, and %d film readers will be in post on 1 April 2027. Under each "
                       "vendor's quote the first contract year would cost, excluding VAT: Halvic Breast %s, Ondrix "
                       "Mammo %s and Lumidra Detect %s. Under protocol 4.4 cost does not change the clinical decision."
                       % ("{:,}".format(f["studies_total"]),
                          "; ".join("%s %s" % (SITE_NAMES[c].split(" Breast")[0], "{:,}".format(f["studies_per_site"][c]))
                                    for c in K.SITE_ORDER), "{:,}".format(f["episodes"]), f["readers_in_post"],
                          "£{:,}".format(f["cost"][a]), "£{:,}".format(f["cost"][q]),
                          "£{:,}".format(f["cost"][v])), D.BODY),
             Paragraph("Counting notes: records the legacy system's history load added for studies already registered "
                       "in the new system during each site's parallel running are counted once; mobile unit 2's "
                       "studies count to Corrowby Road until its vehicle was replaced in May 2026 and to Wendlow "
                       "after; a technical recall booked under the routine code belongs to its woman's episode, "
                       "counted with its routine study; one film reader holds two contracts and is counted once. "
                       "Figures and per-reader results are in ai_tool_evaluation_2027.xlsx; the intervals are charted "
                       "in reader_study_gains.png.", D.NOTE)]
    doc = D._doc(path, "AI reading support for 2027", lead)
    fr = D._frame(K.ORG, "AI reading support: decision paper")
    doc.build(story, onFirstPage=fr, onLaterPages=fr)
