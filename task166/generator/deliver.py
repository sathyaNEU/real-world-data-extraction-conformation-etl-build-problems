"""The goldens, computed from the shipped pack alone (calc.py on target/): onboarding_pod_2027.docx, the decision memo
to the VP Customer Onboarding with its chart, and onboarding_staffing_2025.xlsx, the staffing figures; and figures.json, every
number the write-up states.

    python3 deliver.py <target_dir> <golden_dir> <figures.json>"""
from __future__ import annotations

import datetime as dt
import io
import json
import os
import sys
import warnings

import numpy as np

warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import calc as C          # noqa: E402
import common as K        # noqa: E402
import postprocess as PP  # noqa: E402

F_MEMO = "onboarding_pod_2027.docx"
F_BOOK = "onboarding_staffing_2025.xlsx"
PLANS = ("Core", "Pro", "Enterprise")


def r100(x):
    return int(round(x / 100.0)) * 100


def r200(x):
    return int(round(x / 200.0)) * 200


KINDS = ("one-to-one", "group")


def partner_bookings(T, D):
    import csv
    typ = {r["Account ID"]: r["Type"] for r in D["crm"]}
    return sum(1 for r in csv.DictReader(open(os.path.join(T, C.F["sessions"]))) if typ.get(r["crm_account_id"]) == "Partner")


def compute(T):
    D = C.load(T)
    sp = C.spend(T)
    Y = C.directory(D)
    g, info = C.golden_routes(D, Y, sp)
    order = [int(i) for i in np.argsort(-g)]
    win, second = order[0], order[1]
    dep = Y["dep"]
    sw_wave = sum(dep[win].get(t, 0) for t in range(87, 93))
    sw_25 = sum(dep[win].get(t, 0) for t in range(72, 84)) / 12.0
    lags = [l["lag"] for l in Y["links"] if l["land_t"] <= 92]
    hours = C.ask_pod_hours(T, "golden", D)
    coh = C.ask_cohort(T, "golden", D)
    lag_days = [l["days"] for l in Y["links"] if l["land_t"] <= 92]
    share25 = float(Y["founders"][:, 48 + 72:48 + 84].sum() / D["new"][:, 48 + 72:48 + 84].sum())
    # the usual intake across the golden family and the round 1 judge's constructions on the mobile link
    orgs = [C.golden_routes(D, Y, sp, *f)[1]["org"][win] for f in C.FAMILY]
    orgs += [C.judge_route(D, Y, sp, lf, os_, rt, ls, parts=True)[1][win] for lf in (-48, 48) for os_ in (1, 24, 48)
             for rt in ("pooled", "mature") for ls in (0, -1)]
    fig = dict(
        call=dict(region=C.REG[win], second=C.REG[second], figure=round(float(g[win]), 1), figure_rounded=r200(g[win])),
        components=dict(removals_apr_sep_2026=int(sw_wave), removals_apr_sep_2026_rounded=int(round(sw_wave, -2)),
                        removals_2025_monthly=round(sw_25, 1), removals_ratio=round(sw_wave / 6.0 / sw_25, 2),
                        removals_ratio_rounded=int(round(sw_wave / 6.0 / sw_25)),
                        founder_led_2027=round(float(info["champ"][win]), 1), founder_led_2027_rounded=r100(info["champ"][win]),
                        organic_2027=round(float(info["org"][win]), 1), organic_2027_rounded=r100(info["org"][win]),
                        organic_2027_low=int(round(float(min(orgs)), -1)), organic_2027_high=int(round(float(max(orgs)), -1)),
                        lag_median_months=int(np.median(lags)), lag_median_months_by_days=round(float(np.median(lag_days)) / 30.4375, 1),
                        lag_median_whole_months=int(np.median([d_ // 30.4375 for d_ in lag_days])),
                        lag_median_days=int(np.median(lag_days)),
                        founder_share_2025_pct=int(round(100 * share25))),
        forecast={C.REG[r]: dict(organic=round(float(info["org"][r]), 1), founder_led=round(float(info["champ"][r]), 1),
                                 total=round(float(g[r]), 1)) for r in range(6)},
        hours={"%s Q%d %s" % (p, q, k): round(hours.get((p, q, k), 0.0), 2) for p in C.PODS for q in range(1, 5) for k in KINDS},
        cohort_counts={"%s %s" % (r, p): int(coh["counts"].get((r, p), 0)) for r in C.REG for p in PLANS},
        cohort_acv_k={"%s %s" % (r, p): round(coh["acv"].get((r, p), 0.0) / 1000.0, 1) for r in C.REG for p in PLANS},
        cohort_acv_usd={"%s %s" % (r, p): round(coh["acv"].get((r, p), 0.0), 2) for r in C.REG for p in PLANS},
        cohort_n=int(coh["n"]),
        steps=dict(new_2025=int(D["new"][:, 48 + 72:48 + 84].sum()), links=int(len(Y["links"])),
                   yield_per_100_removals=round(float(np.mean(info["p"])) * 100, 1),
                   same_region_pct=int(round(100 * float(np.mean([l["dest"] == l["origin"] for l in Y["links"]])))),
                   partner_bookings=partner_bookings(T, D), offer_starts=offer_starts(T, D)),
    )
    return fig


def offer_starts(T, D):
    """the 2025 starts whose first-year price fits a plan other than the one their SKU names"""
    A = C.amendments(T)
    launch, _ = C.launch_order_map(D, A, T)
    fit = C.price_plans(T)
    return sum(1 for e in C.cohort_2025(D) for t in [C.start_terms(e, A)] if launch[t[0]] not in fit(t[1], t[2]))


def chart_png(fig):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    for f in ("/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf", "/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf"):
        if os.path.exists(f):
            font_manager.fontManager.addfont(f)
    plt.rcParams.update({"font.family": "Carlito", "font.size": 9.5, "svg.hashsalt": "task166", "axes.edgecolor": "#c3c2b7"})
    ink, ink2, muted = "#0b0b0b", "#52514e", "#898781"
    usual_c, founded_c = "#2a78d6", "#eb6834"
    regs = sorted(C.REG, key=lambda r: fig["forecast"][r]["total"])
    usual = [fig["forecast"][r]["organic"] for r in regs]
    founded = [fig["forecast"][r]["founder_led"] for r in regs]
    total = [a + b for a, b in zip(usual, founded)]
    win = fig["call"]["region"]
    comp = fig["components"]
    f, ax = plt.subplots(figsize=(5.8, 2.75), dpi=200)
    ax.barh(regs, usual, color=usual_c, height=0.56, edgecolor="white", linewidth=1.2, label="Usual intake")
    ax.barh(regs, founded, left=usual, color=founded_c, height=0.56, edgecolor="white", linewidth=1.2,
            label="Founded by people who left paying customers")
    y = regs.index(win)
    ax.text(usual[y] / 2, y, "about %d" % comp["organic_2027_rounded"], va="center", ha="center", fontsize=8.5, color="white")
    ax.text(usual[y] + founded[y] / 2, y, "about %d" % comp["founder_led_2027_rounded"], va="center", ha="center", fontsize=8.5,
            color="white")
    ax.text(total[y] + 8, y, "about %d" % fig["call"]["figure_rounded"], va="center", fontsize=9.5, color=ink, fontweight="bold")
    ax.set_xlim(0, max(total) * 1.18)
    ax.set_xlabel("Expected new paying customers, calendar 2027", color=ink2)
    ax.spines[["top", "right"]].set_visible(False)
    ax.tick_params(axis="y", length=0, labelcolor=ink)
    ax.tick_params(axis="x", labelcolor=muted, color="#c3c2b7")
    for t in ax.get_yticklabels():
        if t.get_text() == win:
            t.set_fontweight("bold")
    ax.grid(axis="x", color="#e1e0d9", linewidth=0.8)
    ax.set_axisbelow(True)
    ax.legend(loc="lower right", frameon=False, fontsize=8, labelcolor=ink2, handlelength=1.2, borderaxespad=0.2)
    ax.set_title("Leavers' new companies put the %s first in 2027" % win, loc="left", fontsize=10.5, fontweight="bold", color=ink)
    f.text(0.01, 0.01, "Source: billing exports, CRM and user directory of 1 October 2026; FY2028 planning basis.", fontsize=7.2,
           color=ink2)
    f.tight_layout(rect=(0, 0.05, 1, 1))
    buf = io.BytesIO()
    f.savefig(buf, format="png", metadata={"Software": None})
    plt.close(f)
    return buf.getvalue()


def write_memo(path, fig):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    c, comp = fig["call"], fig["components"]
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.95)
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.7)
    foot = sec.footer.paragraphs[0]
    foot.text = "Plinthwise Revenue Analytics  |  Internal"
    foot.runs[0].font.size = Pt(8)
    foot.runs[0].font.color.rgb = RGBColor(0x5d, 0x64, 0x6c)
    for line in ("To: %s, VP Customer Onboarding" % K.PEOPLE["vp"], "From: Revenue Analytics", "Date: 23 October 2026",
                 "Re: Where the 2027 onboarding pod goes"):
        p = d.add_paragraph(line)
        p.paragraph_format.space_after = Pt(0)
    h = d.add_paragraph()
    h.paragraph_format.space_before = Pt(10)
    run = h.add_run("Put the 2027 onboarding pod with the %s team" % c["region"])
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x2c, 0x4a, 0x6e)
    d.add_paragraph(
        "The %s is expected to add about %d new paying customers in calendar 2027, more than any other region, and %s comes "
        "second. The pod should sit with the %s team for the year." % (c["region"], c["figure_rounded"], c["second"], c["region"]))
    d.add_paragraph(
        "The %s does not lead on its own history. Fitted on each region's past new customers, trend, season and marketing "
        "plan, the forecast puts %s first and the %s second. What changes the order is where new customers come from: about "
        "two in five of our 2025 new customers were set up by someone who had used Plinthwise at a paying customer, left, "
        "and brought it to a contractor that was not yet one. The new company's first paid invoice comes a median of %d to "
        "%d months after they leave." % (c["region"], c["second"], c["region"], comp["lag_median_whole_months"],
                                         comp["lag_median_months"]))
    d.add_paragraph(
        "Between April and September 2026 the %s's paying customers removed about %s users from their workspaces, about four "
        "times the region's 2025 monthly pace. Most were project managers, superintendents and engineers, and about half came "
        "from the region's largest contractors. At the rate past leavers have founded new customers, and counting each new "
        "company where it bills (close to nine in ten stay in the region they left), they account for about %d of the %s's 2027 new "
        "customers, beside about %d from its usual intake."
        % (c["region"], "{:,}".format(comp["removals_apr_sep_2026_rounded"]), comp["founder_led_2027_rounded"], c["region"],
           comp["organic_2027_rounded"]))
    d.add_paragraph(
        "The forecast is on FP&A's FY2028 planning basis: marketing spend at Marketing Operations' plan, and users removed from "
        "paying customers' workspaces at each region's 2025 monthly average from October 2026. That is the cautious case for "
        "the %s: at the spring and summer pace its lead would widen, and with no removals at all after September it would "
        "still lead." % c["region"])
    d.add_paragraph().add_run().add_picture(io.BytesIO(chart_png(fig)), width=Inches(5.6))
    p = d.add_paragraph()
    r_ = p.add_run("New paying customers as your memo of 20 October counts them: a first paid invoice, or a return more than "
                   "12 months after cancelling, dated in 2027, each company once at its parent, with voided invoices out; "
                   "the region is the parent company's billing state on the March 2024 map. The 2027 figure is rounded to "
                   "the nearest 200 and its parts to the nearest 100.")
    r_.italic = True
    r_.font.size = Pt(8.5)
    d.add_paragraph("The pods' specialist hours and the 2025 customers by starting plan are in onboarding_staffing_2025.xlsx.")
    d.core_properties.author = "Revenue Analytics"
    d.save(path)
    PP.normalize_docx(path, "Revenue Analytics", "Revenue Analytics", "2026-10-22T21:10:00Z", "2026-10-23T16:42:00Z", 4, 96,
                      seed=1661, title="Where the 2027 onboarding pod goes", company=K.COMPANY_LEGAL)


def write_book(path, fig):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    head = PatternFill("solid", fgColor="DCE3EA")
    thin = Side(style="thin", color="9AA3AB")
    ws = wb.active
    ws.title = "Specialist hours 2025"
    ws.append(["Hours of specialist time each pod's specialists spent in sessions held with customers, 2025"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(["Pod", "Sessions", "Q1", "Q2", "Q3", "Q4", "2025"])
    for c in ws[3]:
        c.font = bold
        c.fill = head
    for p in C.PODS:
        for k in KINDS:
            vals = [fig["hours"]["%s Q%d %s" % (p, q, k)] for q in range(1, 5)]
            ws.append([p, k.capitalize(), *vals, round(sum(vals), 2)])
    tot = [round(sum(fig["hours"]["%s Q%d %s" % (p, q, k)] for p in C.PODS for k in KINDS), 2) for q in range(1, 5)]
    ws.append(["All pods", "", *tot, round(sum(tot), 2)])
    for c in ws[ws.max_row]:
        c.font = bold
        c.border = Border(top=thin)
    ws.append([])
    ws.append(["Each specialist's time in a session counts once: the bookings of a field app training or an office hour share "
               "one host and start, and a second host on a booking is that specialist's own hour. A session is held when at "
               "least one of its bookings was completed. Quarter by the session's date. Group sessions are the field app "
               "trainings and office hours."])
    ws.append(["Each session counts for the pod its host was in on the session's date, from the roster's pod assignments; a "
               "session a specialist covered for another pod counts for the host's own pod. A specialist still listed on a "
               "colleague's booking that falls wholly inside a session of their own spent no time on it; the colleague, free "
               "at the time, held it."])
    ws.append(["Source: onboarding_sessions_2024-01_2026-09.csv, onboarding_cases_2024-01_2026-09.csv, "
               "onboarding_team_roster_2026-10-01.csv, CRM accounts."])
    for col, w in zip("ABCDEFG", [12, 12, 9, 9, 9, 9, 10]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=4, max_row=10, min_col=3, max_col=7):
        for c in row:
            c.number_format = "#,##0.00"
            c.alignment = Alignment(horizontal="right")
    ws.freeze_panes = "C4"
    ws2 = wb.create_sheet("2025 new customers")
    ws2.append(["New paying customers of 2025 by the plan they started on, and the annual contract value they signed on it"])
    ws2["A1"].font = Font(bold=True, size=12)
    ws2.append([])
    ws2.append(["Region", "Core", "Pro", "Enterprise", "All plans", "Core ACV ($K)", "Pro ACV ($K)", "Enterprise ACV ($K)"])
    for c in ws2[3]:
        c.font = bold
        c.fill = head
    for r in C.REG:
        n = [fig["cohort_counts"]["%s %s" % (r, p)] for p in PLANS]
        v = [fig["cohort_acv_k"]["%s %s" % (r, p)] for p in PLANS]
        ws2.append([r, *n, sum(n), *v])
    n_all = [sum(fig["cohort_counts"]["%s %s" % (r, p)] for r in C.REG) for p in PLANS]
    v_all = [round(sum(fig["cohort_acv_usd"]["%s %s" % (r, p)] for r in C.REG) / 1000.0, 1) for p in PLANS]
    ws2.append(["All regions", *n_all, sum(n_all), *v_all])
    for c in ws2[ws2.max_row]:
        c.font = bold
        c.border = Border(top=thin)
    ws2.append([])
    ws2.append(["Counted as the 20 October memo counts new paying customers. Plan and ACV in force on the first paid "
                "invoice date: the register's terms today taken back through every amendment effective after that date, a "
                "day's amendments in the order their terms either side chain (a change effective on the first invoice date is "
                "part of the starting terms)."])
    ws2.append(["Plans by SKU as sold: before the catalog of 1 July 2026 the SKUs ran in launch order (PLW-100 Pro, PLW-200 "
                "Enterprise, PLW-300 Core), as the subscriptions' list prices and seats show; the July 2026 catalog "
                "renumbered them in price order. Each customer counts on the plan its SKU names: the PLW-100 subscriptions "
                "first invoiced at Core's list price for their seats are Pro, and each that reached its first anniversary "
                "moved to Pro's price that day."])
    ws2.append(["Source: subscription register (Subscriptions and Amendments), product catalog, legacy billing, CRM."])
    for col, w in zip("ABCDEFGH", [16, 8, 8, 11, 10, 13, 13, 18]):
        ws2.column_dimensions[col].width = w
    for row in ws2.iter_rows(min_row=4, max_row=10, min_col=2, max_col=5):
        for c in row:
            c.number_format = "#,##0"
    for row in ws2.iter_rows(min_row=4, max_row=10, min_col=6, max_col=8):
        for c in row:
            c.number_format = "#,##0.0"
    ws2.freeze_panes = "B4"
    ws3 = wb.create_sheet("Notes")
    for line in ["Prepared by Revenue Analytics for the VP Customer Onboarding, 23 October 2026.",
                 "Exports as pulled on 1 October 2026; activity to 30 September 2026.",
                 "Regions: the sales territory map in effect from 1 March 2024 (Tennessee in Midwest).",
                 "Hours to two decimals; contract value in thousands of dollars to one decimal; the 2027 figure to the nearest 200."]:
        ws3.append([line])
    ws3.column_dimensions["A"].width = 90
    for w_ in (ws, ws2):
        w_.page_setup.orientation = "landscape"
        w_.page_setup.fitToWidth = 1
        w_.page_setup.fitToHeight = 0
        w_.sheet_properties.pageSetUpPr.fitToPage = True
        w_.print_title_rows = "3:3"
    for w_ in wb.worksheets:
        w_.sheet_view.showGridLines = w_.title != "Notes"
    wb.save(path)
    PP.normalize_ooxml(path, "Revenue Analytics", "Revenue Analytics", "2026-10-22T20:05:00Z", "2026-10-23T16:30:00Z",
                       excel=True, company=K.COMPANY_LEGAL)


def main(target, golden, figpath):
    os.makedirs(golden, exist_ok=True)
    fig = compute(target)
    write_memo(os.path.join(golden, F_MEMO), fig)
    write_book(os.path.join(golden, F_BOOK), fig)
    with open(figpath, "w") as fh:
        json.dump(fig, fh, indent=1, sort_keys=False)
        fh.write("\n")
    return fig


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3])
