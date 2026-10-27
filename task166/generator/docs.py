"""task166's documents: the VP Customer Onboarding's memo that files the decision and the count, FP&A's FY2028
planning basis, the FY2023 board deck's customer-growth page, the data dictionary and the export manifest. The text
is kept here as data so checks.py and verify.py can quote it and read it back."""
from __future__ import annotations

import datetime as dt
import os
import re

import numpy as np
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import LETTER, landscape                                 # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import inch                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import common as K          # noqa: E402
import onboarding as OB     # noqa: E402
import records as RC        # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/"
for _n, _f in (("LibSans", "liberation/LiberationSans-Regular.ttf"), ("LibSans-Bold", "liberation/LiberationSans-Bold.ttf")):
    pdfmetrics.registerFont(TTFont(_n, FONT_DIR + _f))
INK = colors.HexColor("#1f2328")
GREY = colors.HexColor("#5d646c")
ACCENT = colors.HexColor("#2b4b6f")
RULE = colors.HexColor("#c3c9cf")
HEAD = colors.HexColor("#e8ecf0")

F_MEMO = "memo_2027_onboarding_pod.docx"
F_BASIS = "fy2028_planning_basis.pdf"
F_DECK = "board_deck_2023-03_customer_growth.pdf"
F_DICT = "data_dictionary.pdf"
F_MANIFEST = "export_manifest.csv"

# ------------------------------------------------------------------------------------------------- the memo
MEMO_HEAD = [("Plinthwise, Customer Onboarding", "h"), ("To: Revenue Analytics", "p"),
             ("From: %s, VP Customer Onboarding" % K.PEOPLE["vp"], "p"), ("Date: 20 October 2026", "p"),
             ("Re: Where the 2027 onboarding pod goes", "p")]
MEMO = [
    "The FY2028 plan funds a fourth onboarding pod beside West, Central and East, which serve the Southwest and Pacific, "
    "Texas & Gulf and the Midwest, and the Southeast and Northeast: four implementation specialists who run the first 90 "
    "days of every new customer. Finance has agreed to open its requisitions in November so the pod is staffed before "
    "January, and it sits with one regional sales team for all of calendar 2027. It goes to the region expected to add the "
    "most new paying customers in calendar 2027, forecast on FP&A's FY2028 planning basis.",
    "A new paying customer is a company that starts paying us in calendar 2027: its first paid subscription invoice is dated "
    "in 2027, or it comes back in 2027 more than 12 months after it cancelled. A company that comes back within 12 months "
    "of cancelling is a reactivation, not a new customer. Count each company once, at its parent: a new division of a "
    "company that already pays us is not a new customer. A voided invoice is not a paid one. Each new customer counts in "
    "the region of its own parent company's billing state, under the territory map that took effect on 1 March 2024, when "
    "Tennessee moved from Southeast to Midwest.",
    "The billing history is in two systems: the legacy export runs to February 2023 and the subscription register starts "
    "with the December 2022 to February 2023 dual run. The CRM export has the account crosswalk, Marketing Operations' "
    "workbook has spend actuals and the plan through January 2028, and Sales Operations' file has the territory map. The "
    "product's user directory, the onboarding team's case and session exports and the team roster are as pulled on 1 "
    "October.",
    "Please send me the region, your 2027 figure for it and the region that comes second, with the staffing figures I asked "
    "for and a short note on how you got there.",
]
MEMO_SIGN = "Jeff"

# ------------------------------------------------------------------------------------------------- the planning basis
BASIS_SCOPE = (
    "These are the assumptions every FY2028 plan line uses, and every calendar-2027 view prepared from the plan. FY2028 "
    "runs from 1 February 2027 to 31 January 2028. A calendar-2027 view takes January 2027 from FY2027 (period 12) and "
    "February to December 2027 from FY2028 (periods 1 to 11).")
BASIS_ROWS = [
    ("Pricing", "List prices are held through 31 January 2027. From 1 February 2027 Core and Pro renew at list plus 4 per "
                "cent; Enterprise renews on its contract terms."),
    ("Marketing programmes", "Marketing Operations' plan by region and month, October 2026 to January 2028, as set out on "
                             "the plan sheet of its spend workbook."),
    ("Sales capacity", "Account executives held at the September 2026 headcount in each region; no new territories."),
    ("Cancellations", "Each region's customers cancel at the region's rate for the twelve months to 30 September 2026."),
    ("Seats", "From October 2026, users are added to and removed from paying customers' workspaces at each region's 2025 "
              "monthly average. A removal is a user deactivated on the workspace of a customer that is still paying; users "
              "deactivated when a customer's account closes belong to the cancellations line."),
    ("Hosting", "$1.85 per active user per month, the rate in the cloud contract renewed in June 2026."),
    ("Collections", "38 days sales outstanding."),
    ("Territories", "The sales territory map in effect from 1 March 2024."),
]
BASIS_CHANGES = ("Changes from the FY2027 basis: the pricing step moves from 1 March to 1 February, and hosting is at the "
                 "June 2026 renewal rate.")

# ------------------------------------------------------------------------------------------------- the dictionary
DICT = [
    (RC.F_LEGACY, "Legacy billing system: every subscription with a first invoice before 1 March 2023.", [
        ("customer_code", "the legacy billing customer"), ("subscription_id", "the legacy subscription"),
        ("company, billing_state", "the customer's name and bill-to state"), ("plan, currency, mrr_cents", "the plan and monthly price"),
        ("subscription_created", "the date the subscription was set up"), ("first_invoice_date", "the date of the subscription's first invoice"),
        ("first_invoice_status", "paid or voided"), ("cancelled_on", "cancellation date, where cancelled in this system"),
        ("notes", "billing notes")]),
    (RC.F_REG, "Subscription register: every subscription with a service period on or after 1 December 2022 (sheet "
               "Subscriptions), and the amendments recorded from that date (sheet Amendments).", [
        ("Account Number", "the billing account"), ("Subscription Number", "the subscription"),
        ("Parent Account Number", "the billing account it rolls up to, where it has one"),
        ("Account Name, Bill To State", "the account's name and bill-to state"), ("Rate Plan SKU, Seats, ACV (USD)", "the rate plan's SKU, seats and annual contract value"),
        ("First Invoice Date, First Invoice Status", "the first invoice's date, and Paid or Voided"),
        ("Service Start", "the first day of service in this system"), ("Subscription Status, Cancellation Date", "Active or Cancelled, and when"),
        ("Last Amended", "the date of the latest amendment"), ("Migrated From (legacy subscription)", "the legacy subscription_id it continues"),
        ("Effective Date, Amendment Number", "Amendments sheet: the date the amendment took effect, and its reference"),
        ("Change", "Amendments sheet: the kind of amendment"),
        ("SKU, Seats and ACV (USD) Before and After", "Amendments sheet: the subscription's terms either side of the amendment")]),
    (RC.F_CATALOG, "Billing platform: the product catalog.", [
        ("sku, plan", "the rate plan SKU and the plan it sells"), ("list_price_month_usd, seats_included, extra_seat_month_usd",
                                                                    "the monthly list price, the seats it includes and the monthly price of each further seat"),
        ("in_catalog_from", "the date the SKU entered this catalog")]),
    (RC.F_CRM, "CRM: customer, former customer and partner accounts at 1 October 2026.", [
        ("Account ID, Account Name", "the CRM account"), ("Parent Account ID", "the parent CRM account, where it has one"),
        ("Billing State", "bill-to state"), ("Sales Region", "the sales region recorded when the account was created"),
        ("Type", "Customer, Former Customer or Partner"), ("Lead Source", "the lead source recorded at creation"),
        ("Created Date", "the date the account was created"), ("Billing Account Number", "the subscription register's Account Number"),
        ("Legacy Customer Code", "the legacy billing customer_code"), ("Product Workspace ID", "the account's product workspace")]),
    (RC.F_DIR, "Product user directory: every user account on every workspace, as exported on 1 October 2026.", [
        ("user_id", "the user account"), ("workspace_id", "the workspace"), ("full_name, email, mobile, job_title", "as the user or an admin entered them"),
        ("workspace_role", "Admin or Member"), ("created_at", "when the user account was created (UTC)"),
        ("last_active_at", "the latest sign-in (UTC)"), ("deactivated_at", "when the user account was deactivated (UTC)"),
        ("deactivation_reason", "the reason recorded at deactivation")]),
    (RC.F_TERR, "Sales Operations: the sales territory map.", [
        ("state, state_name", "the bill-to state"), ("sales_region", "the state's region from 1 March 2024"),
        ("sales_region_before_2024_03_01", "the state's region before then")]),
    (RC.F_SPEND, "Marketing Operations: programme spend in thousands of dollars.", [
        ("Actuals", "fiscal year, period, calendar month, region, channel and spend, February 2019 to September 2026"),
        ("Plan Oct-2026 to Jan-2028", "fiscal year, period, calendar month, region and planned spend")]),
    (OB.F_CASES, "Onboarding tool: onboarding cases opened from January 2024.", [
        ("case_id", "the onboarding case"), ("crm_account_id", "the CRM account onboarded"), ("account_name", "its name"),
        ("playbook", "the onboarding playbook the case follows"), ("opened_on", "the date the case was opened"),
        ("owner", "the specialist who owns the case"), ("rate_plan_sku, seats, contract_acv_usd", "the rate plan SKU, seats and annual contract value on the case"),
        ("go_live_on", "the date of the go-live review"), ("status", "the case's status at 1 October 2026")]),
    (OB.F_SESSIONS, "Scheduling tool: onboarding sessions booked from January 2024.", [
        ("booking_id", "the scheduling tool's booking reference"), ("case_id", "the onboarding case"),
        ("crm_account_id", "the CRM account the booking was made under"),
        ("session_type", "the session template: Kickoff, Admin setup and Go-live review are one-to-one sessions; Field app "
                         "training and Office hours are group sessions, with a booking for each customer who joins"),
        ("host", "a specialist assigned to the booking; a booking with two specialists has a row for each"),
        ("start_utc", "the booked start (UTC)"),
        ("duration_min", "the booked length in minutes"), ("status", "Completed, No-show, Rescheduled or Cancelled by customer"),
        ("attendees", "people from the customer who joined")]),
    (OB.F_ROSTER, "HR: the onboarding team and each specialist's pod assignments.", [
        ("employee_id, name", "the specialist"),
        ("pod, pod_from, pod_to", "an onboarding pod the specialist worked in and the first and last day there (pod_to blank for "
                                  "the current pod); a specialist who moved pod has a row for each"),
        ("title", "the job title"), ("started_on, left_on", "the first and, for someone who has left the team, the last day on it")]),
    (OB.F_CSAT, "Survey tool: post-onboarding survey responses.", [
        ("response_id, case_id", "the response and its case"), ("survey_sent_on, responded_on", "when sent and answered"),
        ("onboarding_score_1_5, likely_to_recommend_0_10", "the two questions")]),
]


def _styles():
    return dict(
        h1=ParagraphStyle("h1", fontName="LibSans-Bold", fontSize=14.5, leading=18.5, textColor=ACCENT, spaceAfter=2),
        sub=ParagraphStyle("sub", fontName="LibSans", fontSize=8.6, leading=11.4, textColor=GREY, spaceAfter=10),
        h2=ParagraphStyle("h2", fontName="LibSans-Bold", fontSize=10.4, leading=14, textColor=INK, spaceBefore=6, spaceAfter=3),
        body=ParagraphStyle("b", fontName="LibSans", fontSize=9.6, leading=13.2, textColor=INK, spaceAfter=5),
        cell=ParagraphStyle("c", fontName="LibSans", fontSize=8.4, leading=10.6, textColor=INK),
        cellb=ParagraphStyle("cb", fontName="LibSans-Bold", fontSize=8.4, leading=10.6, textColor=INK),
        small=ParagraphStyle("sm", fontName="LibSans", fontSize=7.8, leading=10, textColor=GREY, spaceAfter=4))


def _footer(text, page=LETTER):
    def foot(canvas, doc):
        canvas.saveState()
        canvas.setFont("LibSans", 7.4)
        canvas.setFillColor(GREY)
        canvas.drawString(0.9 * inch, 0.55 * inch, text)
        canvas.drawRightString(page[0] - 0.9 * inch, 0.55 * inch, "Page %d" % doc.page)
        canvas.restoreState()
    return foot


def _table(data, widths, head=True, align_right_from=None):
    t = Table(data, colWidths=widths, repeatRows=1 if head else 0)
    style = [("GRID", (0, 0), (-1, -1), 0.4, RULE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("TOPPADDING", (0, 0), (-1, -1), 2.2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2)]
    if head:
        style.append(("BACKGROUND", (0, 0), (-1, 0), HEAD))
    if align_right_from is not None:
        style.append(("ALIGN", (align_right_from, 0), (-1, -1), "RIGHT"))
    t.setStyle(TableStyle(style))
    return t


def write_memo(path):
    import docx
    from docx.shared import Pt
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for text, kind in MEMO_HEAD:
        if kind == "h":
            d.add_heading(text, level=2)
        else:
            d.add_paragraph(text)
    for p in MEMO:
        d.add_paragraph(p)
    d.add_paragraph(MEMO_SIGN)
    d.core_properties.author = K.PEOPLE["vp"]
    d.save(path)


def write_basis(path):
    st = _styles()
    s = [Paragraph("FY2028 planning basis", st["h1"]),
         Paragraph("%s  |  Finance (FP&amp;A)  |  Owner: %s, Director of FP&amp;A  |  Version 1.0, issued 2 October 2026"
                   % (K.COMPANY_LEGAL, K.PEOPLE["fpa"]), st["sub"]),
         Paragraph("What this covers", st["h2"]), Paragraph(BASIS_SCOPE, st["body"]),
         Paragraph("Assumptions", st["h2"])]
    data = [[Paragraph("Area", st["cellb"]), Paragraph("Assumption", st["cellb"])]]
    for a, b in BASIS_ROWS:
        data.append([Paragraph(a, st["cell"]), Paragraph(b.replace("&", "&amp;"), st["cell"])])
    s.append(_table(data, [1.45 * inch, 5.25 * inch]))
    s.append(Spacer(1, 8))
    s.append(Paragraph(BASIS_CHANGES, st["small"]))
    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=0.9 * inch, rightMargin=0.9 * inch, topMargin=0.8 * inch,
                            bottomMargin=0.85 * inch, title="FY2028 planning basis", author=K.PEOPLE["fpa"], subject="",
                            creator="", producer="")
    foot = _footer("%s  |  FP&A  |  FY2028 planning basis  |  v1.0" % K.COMPANY)
    doc.build(s, onFirstPage=foot, onLaterPages=foot)


def deck_figures(B):
    """New customers by fiscal year (FY2020 to FY2023) and region as assigned at the time, and ARR at 31 January 2023."""
    out = {r: [0, 0, 0, 0] for r in K.REGIONS}
    arr = {r: 0.0 for r in K.REGIONS}
    for c in B["comps"]:
        fi = c["first_inv"]
        fy = fi.year + 1 if fi.month >= 2 else fi.year
        if 2020 <= fy <= 2023:
            reg = K.region_of_state(c["state"], fi)
            out[reg][fy - 2020] += 1
        end = dt.date(2023, 1, 31)
        if fi <= end and (c["cancel"] is None or c["cancel"] > end):
            reg = K.region_of_state(c["state"], end)
            arr[reg] += RC._state_at(c, end)[2]
            for dv in c["divisions"]:
                if dv["first_inv"] <= end and (dv["cancel"] is None or dv["cancel"] > end):
                    arr[reg] += dv["acv"]
    return out, arr


def write_deck(path, B):
    st = _styles()
    page = landscape(LETTER)
    new, arr = deck_figures(B)
    s = [Paragraph("Customer growth, FY2020 to FY2023", st["h1"]),
         Paragraph("Board of Directors, 14 March 2023  |  Section 3 of 7  |  Prepared by Revenue Operations", st["sub"])]
    data = [["Sales region", "FY2020", "FY2021", "FY2022", "FY2023", "ARR at 31 Jan 2023 ($M)"]]
    for r in K.REGIONS:
        data.append([r] + ["{:,}".format(v) for v in new[r]] + ["%.1f" % (arr[r] / 1e6)])
    tot = [sum(new[r][i] for r in K.REGIONS) for i in range(4)]
    data.append(["Total"] + ["{:,}".format(v) for v in tot] + ["%.1f" % (sum(arr.values()) / 1e6)])
    s.append(_table(data, [1.9 * inch] + [1.05 * inch] * 4 + [1.9 * inch], align_right_from=1))
    s.append(Spacer(1, 8))
    s.append(Paragraph("New customers: companies whose first paid invoice fell in the fiscal year, counted at the parent; "
                       "reactivations within 12 months excluded. Regions as assigned at the time. ARR: annual contract value "
                       "of subscriptions in service at 31 January 2023, divisions included.", st["small"]))
    doc = SimpleDocTemplate(path, pagesize=page, leftMargin=0.9 * inch, rightMargin=0.9 * inch, topMargin=0.8 * inch,
                            bottomMargin=0.85 * inch, title="Board deck, March 2023: customer growth", author="Revenue Operations",
                            subject="", creator="", producer="")
    foot = _footer("%s  |  Board of Directors  |  March 2023  |  Confidential" % K.COMPANY, page)
    doc.build(s, onFirstPage=foot, onLaterPages=foot)
    return new, arr


def write_dictionary(path):
    st = _styles()
    s = [Paragraph("Data dictionary: exports of 1 October 2026", st["h1"]),
         Paragraph("%s  |  Revenue Analytics  |  The files pulled for the 2027 onboarding pod request" % K.COMPANY_LEGAL, st["sub"])]
    for f, what, cols in DICT:
        s.append(Paragraph(f, st["h2"]))
        s.append(Paragraph(what, st["body"]))
        data = [[Paragraph("Field", st["cellb"]), Paragraph("Meaning", st["cellb"])]]
        for a, b in cols:
            data.append([Paragraph(a, st["cell"]), Paragraph(b, st["cell"])])
        s.append(_table(data, [2.3 * inch, 4.4 * inch]))
    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=0.9 * inch, rightMargin=0.9 * inch, topMargin=0.8 * inch,
                            bottomMargin=0.85 * inch, title="Data dictionary", author="Revenue Analytics", subject="", creator="",
                            producer="")
    foot = _footer("%s  |  Revenue Analytics  |  Data dictionary" % K.COMPANY)
    doc.build(s, onFirstPage=foot, onLaterPages=foot)


MANIFEST = [
    (RC.F_LEGACY, "Legacy billing", "Revenue Operations", "2026-10-01T06:40:00-06:00", "first invoices January 2015 to February 2023"),
    (RC.F_REG, "Subscription register", "Revenue Operations", "2026-10-01T07:15:00-06:00", "service periods and amendments from 1 December 2022"),
    (RC.F_CRM, "CRM", "Sales Operations", "2026-10-01T07:02:00-06:00", "customer, former customer and partner accounts"),
    (RC.F_DIR, "Product (user directory)", "Data Engineering", "2026-10-01T05:30:00-06:00", "every user account to 30 September 2026"),
    (RC.F_TERR, "Sales Operations", "Sales Operations", "2026-10-01T08:10:00-06:00", "the map from 1 March 2024 and the one before"),
    (RC.F_SPEND, "Marketing Operations", "Marketing Operations", "2026-10-01T09:45:00-06:00", "actuals February 2019 to September 2026; plan October 2026 to January 2028"),
    (OB.F_CASES, "Onboarding tool", "Customer Onboarding", "2026-10-01T08:30:00-06:00", "cases opened January 2024 to September 2026"),
    (OB.F_SESSIONS, "Scheduling tool", "Customer Onboarding", "2026-10-01T08:31:00-06:00", "bookings January 2024 to September 2026"),
    (OB.F_CSAT, "Survey tool", "Customer Onboarding", "2026-10-01T08:36:00-06:00", "responses January 2024 to September 2026"),
    (OB.F_ROSTER, "HR", "People Operations", "2026-10-01T09:05:00-06:00", "the onboarding team's pod assignments at 1 October 2026, with the specialists who left it since 2024"),
    (RC.F_CATALOG, "Billing platform", "Revenue Operations", "2026-10-01T07:16:00-06:00", "the product catalog"),
    (F_MEMO, "Customer Onboarding", "Customer Onboarding", "2026-10-20T16:05:00-06:00", "the request"),
    (F_BASIS, "FP&A", "Finance", "2026-10-02T11:00:00-06:00", "FY2028 planning basis v1.0"),
    (F_DECK, "Board materials", "Revenue Operations", "2023-03-14T09:00:00-06:00", "FY2023 board deck, customer growth page"),
    (F_DICT, "Revenue Analytics", "Revenue Analytics", "2026-10-02T14:20:00-06:00", "field definitions"),
]


def write_manifest(path, n):
    rows = []
    for f, sysname, team, when, covers in MANIFEST:
        rows.append([f, sysname, team, when, n.get(f, ""), covers])
    rows.append([F_MANIFEST, "Revenue Analytics", "Revenue Analytics", "2026-10-02T14:25:00-06:00", "", "this list"])
    return K.write_csv(path, ["file", "source", "owner", "exported_at", "rows", "covers"], rows)


def write_all(out, B, n):
    write_memo(os.path.join(out, F_MEMO))
    write_basis(os.path.join(out, F_BASIS))
    B["deck"] = write_deck(os.path.join(out, F_DECK), B)
    write_dictionary(os.path.join(out, F_DICT))
    n[F_MANIFEST] = write_manifest(os.path.join(out, F_MANIFEST), n)
    return n
