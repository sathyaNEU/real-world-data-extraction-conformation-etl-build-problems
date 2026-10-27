"""task166's independent verifier. It reads only the shipped bundle (target/, golden/, submission.md) on its own code path
(the csv module, numpy, openpyxl, python-docx and pypdf; no generator module and no pandas), recounts the new paying
customers the memo's way, links each new workspace's first user to the people removed from paying customers, rebuilds
the founder pipeline and the organic count model, recomputes the call, its figure, the region second and all 60 ask
figures, reads back every format the pack ships, and compares the goldens, figures.json and submission.md with what it
recomputed.

    python3 task166/generator/verify.py [target]"""
from __future__ import annotations

import csv
import datetime as dt
import json
import math
import os
import re
import sys
from collections import defaultdict

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
TARGET = sys.argv[1] if len(sys.argv) > 1 else os.path.join(TASK, "target")
GOLDEN = os.path.join(TASK, "golden")
REGIONS = ["Southwest", "Southeast", "Texas & Gulf", "Midwest", "Northeast", "Pacific"]
NICKNAMES = {"Bob": "Robert", "Bill": "William", "Rick": "Richard", "Joe": "Joseph", "Tom": "Thomas", "Chris": "Christopher",
             "Dan": "Daniel", "Matt": "Matthew", "Tony": "Anthony", "Mike": "Michael", "Jim": "James", "Steve": "Steven",
             "Andy": "Andrew", "Ken": "Kenneth", "Tim": "Timothy", "Ron": "Ronald", "Ed": "Edward", "Jeff": "Jeffrey",
             "Nick": "Nicholas", "Jon": "Jonathan", "Greg": "Gregory", "Alex": "Alexander", "Pat": "Patrick",
             "Liz": "Elizabeth", "Jen": "Jennifer", "Kate": "Katherine", "Becky": "Rebecca", "Sam": "Samuel",
             "Debbie": "Deborah", "Kim": "Kimberly", "Ben": "Benjamin", "Zach": "Zachary"}
FAILS = []


def ok(label, cond, msg):
    print("%s  %-5s %s" % ("PASS" if cond else "FAIL", label, msg))
    if not cond:
        FAILS.append(label)


def rows(name):
    with open(os.path.join(TARGET, name), newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def day(v):
    if v in (None, ""):
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return dt.date(int(str(v)[:4]), int(str(v)[5:7]), int(str(v)[8:10]))


def stamp(v):
    return dt.datetime(int(v[:4]), int(v[5:7]), int(v[8:10]), int(v[11:13]), int(v[14:16]), int(v[17:19])) if v else None


def mon(d):
    """months since January 2019"""
    return 12 * (d.year - 2019) + d.month - 1


def sheet(path, title):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = [list(r) for r in wb[title].iter_rows(values_only=True)]
    wb.close()
    return out


# ------------------------------------------------------------------------------------------------- formats read back
files = sorted(os.listdir(TARGET))
fmts = sorted({f.rsplit(".", 1)[1] for f in files})
ok("V1", len(files) >= 10 and len(fmts) >= 3, "%d files in %d formats (%s)" % (len(files), len(fmts), ", ".join(fmts)))
csvs = {f: rows(f) for f in files if f.endswith(".csv")}
ok("V2", all(len(v) > 0 for v in csvs.values()) and max(len(v) for v in csvs.values()) >= 25000,
   "every csv reads; the largest has %d rows" % max(len(v) for v in csvs.values()))
import docx                 # noqa: E402
import pypdf                # noqa: E402
texts = {}
for f in files:
    if f.endswith(".pdf"):
        texts[f] = " ".join((p.extract_text() or "") for p in pypdf.PdfReader(os.path.join(TARGET, f)).pages)
    elif f.endswith(".docx"):
        texts[f] = " ".join(p.text for p in docx.Document(os.path.join(TARGET, f)).paragraphs)
ok("V3", all(len(t) > 400 for t in texts.values()), "every pdf and docx has a text layer (%s)" % ", ".join(sorted(texts)))
F_REG = [f for f in files if f.startswith("billing_subscription_register")][0]
F_SPEND = [f for f in files if f.startswith("marketing_spend")][0]
subs_sheet = sheet(os.path.join(TARGET, F_REG), "Subscriptions")
amend_sheet = sheet(os.path.join(TARGET, F_REG), "Amendments")
ok("V4", len(subs_sheet) > 10000 and len(amend_sheet) > 1000, "the register reads: %d subscriptions, %d amendments"
   % (len(subs_sheet) - 2, len(amend_sheet) - 2))

# ------------------------------------------------------------------------------------------------- the memo's count
region_of_state = {r["state"]: REGIONS.index(r["sales_region"]) for r in csvs["sales_territories_2024-03.csv"]}
crm = csvs["crm_accounts_2026-10-01.csv"]
acct = {r["Account ID"]: r for r in crm}
by_code = {r["Legacy Customer Code"]: r["Account ID"] for r in crm if r["Legacy Customer Code"]}
by_bill = {r["Billing Account Number"]: r["Account ID"] for r in crm if r["Billing Account Number"]}


def top(a):
    for _ in range(6):
        if not acct[a]["Parent Account ID"]:
            break
        a = acct[a]["Parent Account ID"]
    return a


S = {}
for r in csvs["billing_legacy_subscriptions_2015-2023.csv"]:
    S[("L", r["subscription_id"])] = dict(no=r["subscription_id"], crm=by_code.get(r["customer_code"]), first=day(r["first_invoice_date"]),
                                          paid=r["first_invoice_status"] == "paid", cancel=day(r["cancelled_on"]), reg=None)
hdr = subs_sheet[1]
for v in subs_sheet[2:]:
    r = dict(zip(hdr, v))
    if not r["Subscription Number"]:
        continue
    rec = dict(no=r["Subscription Number"], crm=by_bill.get(r["Account Number"]), first=day(r["First Invoice Date"]),
               paid=r["First Invoice Status"] == "Paid", cancel=day(r["Cancellation Date"]), plan=r["Rate Plan SKU"],
               seats=int(r["Seats"]), acv=float(r["ACV (USD)"]), reg=r["Subscription Number"])
    m = r["Migrated From (legacy subscription)"]
    if m and ("L", m) in S:
        L = S[("L", m)]
        L["cancel"] = L["cancel"] or rec["cancel"]
        L["crm"] = L["crm"] or rec["crm"]
        L["reg"] = rec["no"]
        L["plan"], L["seats"], L["acv"] = rec["plan"], rec["seats"], rec["acv"]
        L["plan"], L["acv"] = rec["plan"], rec["acv"]
        continue
    S[("R", rec["no"])] = rec
paid = [s for s in S.values() if s["paid"] and s["crm"]]
company = defaultdict(list)
for s in paid:
    s["top"] = top(s["crm"])
    company[s["top"]].append(s)
events = []
for c, ss in company.items():
    own = sorted([s for s in ss if s["crm"] == c], key=lambda s: (s["first"], s["no"]))
    for i, s in enumerate(own):
        if i:
            prior = own[:i]
            if any(p["cancel"] is None or p["cancel"] >= s["first"] for p in prior):
                continue
            if (s["first"] - max(p["cancel"] for p in prior)).days <= 366:
                continue
        st = acct[c]["Billing State"]
        if st in region_of_state and dt.date(2015, 1, 1) <= s["first"] <= dt.date(2026, 9, 30):
            events.append(dict(top=c, first=s["first"], region=region_of_state[st], sub=s, t=mon(s["first"]),
                               ws=acct[c]["Product Workspace ID"]))
counts = np.zeros((6, 156))          # months from January 2015 (t = -48) to December 2027
for e in events:
    counts[e["region"], e["t"] + 48] += 1
n25 = int(sum(1 for e in events if e["first"].year == 2025))
ok("V5", 1500 <= n25 <= 2000, "the memo's count: %d new paying customers in 2025, %d events since 2015" % (n25, len(events)))

# ------------------------------------------------------------------------------------------------- spend
act = defaultdict(float)
for v in sheet(os.path.join(TARGET, F_SPEND), "Actuals")[1:]:
    y, m = map(int, v[2].split("-"))
    act[(REGIONS.index(v[3]), 12 * (y - 2019) + m - 1)] += float(v[5])
for v in sheet(os.path.join(TARGET, F_SPEND), "Plan Oct-2026 to Jan-2028")[1:]:
    y, m = map(int, v[2].split("-"))
    act[(REGIONS.index(v[3]), 12 * (y - 2019) + m - 1)] = float(v[4])


def poisson_fit(y, X):
    """Newton-Raphson on the Poisson log-likelihood, with step halving."""
    b = np.zeros(X.shape[1])
    b[0] = math.log(max(y.mean(), 1e-6))

    def ll(beta):
        eta = X @ beta
        return float(np.sum(y * eta - np.exp(eta)))
    cur = ll(b)
    for _ in range(200):
        mu = np.exp(X @ b)
        g = X.T @ (y - mu)
        H = (X * mu[:, None]).T @ X
        step = np.linalg.solve(H, g)
        k = 1.0
        while k > 1e-6:
            nb = b + k * step
            nl = ll(nb)
            if nl >= cur - 1e-12:
                break
            k /= 2
        if np.max(np.abs(nb - b)) < 1e-11:
            b = nb
            break
        b, cur = nb, nl
    return b


def design(t, r, sbar, with_spend=True, region_trend=True):
    x = [1.0 if j == r else 0.0 for j in range(6)]
    x += [t / 12.0 if (j == r or not region_trend) else 0.0 for j in (range(6) if region_trend else [0])]
    x += [1.0 if (t % 12) == j else 0.0 for j in range(1, 12)]
    if with_spend:
        x.append(math.log(act[(r, t)] / sbar[r]))
    return x


def count_model(series, with_spend=True, region_trend=True, t0=1, t1=92):
    ts = range(t0, t1 + 1)
    sbar = [float(np.mean([act[(r, t)] for t in ts])) for r in range(6)]
    X = np.array([design(t, r, sbar, with_spend, region_trend) for r in range(6) for t in ts])
    y = np.array([series[r, t + 48] for r in range(6) for t in ts], float)
    b = poisson_fit(y, X)
    return np.array([sum(math.exp(np.dot(design(t, r, sbar, with_spend, region_trend), b)) for t in range(96, 108))
                     for r in range(6)])


rung0 = count_model(counts)
ok("V6", int(np.argmax(rung0)) == 2, "rung 0 (region trend, season and spend on each region's own new customers): %s"
   % ", ".join("%s %.1f" % (REGIONS[j], rung0[j]) for j in np.argsort(-rung0)[:3]))

# ------------------------------------------------------------------------------------------------- the directory
users = csvs["product_user_directory_2026-10-01.csv"]
ws_top = {r["Product Workspace ID"]: r["Account ID"] for r in crm if r["Product Workspace ID"] and not r["Parent Account ID"]}
first_paid = {}
cancels = defaultdict(list)
for c, ss in company.items():
    own = [s["first"] for s in ss if s["crm"] == c]
    if own:
        first_paid[c] = min(own)
    cancels[c] = [s["cancel"] for s in ss if s["cancel"]]


def digits(p):
    d = "".join(ch for ch in p if ch.isdigit())
    return d[-10:] if len(d) >= 10 else ""


def person(n):
    parts = [w for w in re.sub(r"[^A-Za-z ]", " ", n).split() if len(w) > 1]
    return ("%s %s" % (NICKNAMES.get(parts[0], parts[0]), parts[-1])).lower() if parts else ""


removed = []
for u in users:
    if not u["deactivated_at"]:
        continue
    c = ws_top.get(u["workspace_id"])
    if c is None or acct[c]["Billing State"] not in region_of_state:
        continue
    when = stamp(u["deactivated_at"])
    local = (when - dt.timedelta(hours=7)).date()
    if c in first_paid and local < first_paid[c] - dt.timedelta(days=60):
        continue
    if u["deactivation_reason"] == "Account closed" or any(abs((when.date() - x).days) <= 3 for x in cancels[c]):
        continue
    removed.append(dict(ws=u["workspace_id"], region=region_of_state[acct[c]["Billing State"]], d=local, t=mon(local),
                        phone=digits(u["mobile"]), name=person(u["full_name"])))
dep = np.zeros((6, 160))
for x in removed:
    if x["t"] >= -48:
        dep[x["region"], x["t"] + 48] += 1
sw_wave = int(dep[0, 87 + 48:93 + 48].sum())
sw_2025 = dep[0, 72 + 48:84 + 48].sum() / 12
ok("V7", 6000 <= sw_wave <= 7000, "Southwest removals April to September 2026: %d, %.2f times its 2025 monthly %.1f"
   % (sw_wave, sw_wave / 6 / sw_2025, sw_2025))
phone_ix, name_ix = defaultdict(list), defaultdict(list)
for x in removed:
    if x["phone"]:
        phone_ix[x["phone"]].append(x)
    if x["name"]:
        name_ix[x["name"]].append(x)
on_ws = defaultdict(list)
for u in users:
    on_ws[u["workspace_id"]].append(u)
links = []
for e in events:
    us = sorted(on_ws.get(e["ws"], []), key=lambda u: (u["created_at"], u["user_id"]))
    if not us:
        continue
    f = us[0]
    k = digits(f["mobile"])
    if k:
        cand = [x for x in phone_ix.get(k, []) if x["ws"] != e["ws"]]
    else:
        cand = [x for x in name_ix.get(person(f["full_name"]), []) if x["ws"] != e["ws"] and x["region"] == e["region"]]
    cand = [x for x in cand if 0 < (e["first"] - x["d"]).days <= 1116]
    if cand:
        x = max(cand, key=lambda z: z["d"])
        lag = max(1, 12 * (e["first"].year - x["d"].year) + e["first"].month - x["d"].month)
        links.append(dict(lag=lag, land=e["t"], dest=e["region"], origin=x["region"]))
pmf = np.bincount(np.array([l["lag"] for l in links if l["land"] <= 92]) - 1, minlength=40)[:40].astype(float)
pmf /= pmf.sum()
tot_dep = dep.sum(0)
expo = sum(tot_dep[t + 48] * pmf[max(0, -36 - t - 1):max(0, 92 - t)].sum() for t in range(-48, 93) if 92 - t > 0)
yield_ = sum(1 for l in links if -36 <= l["land"] <= 92) / expo
same = np.mean([l["dest"] == l["origin"] for l in links])
share = np.full((6, 6), (1 - same) / 5)
np.fill_diagonal(share, same)
future = dep[:, 72 + 48:84 + 48].sum(1) / 12
founders27 = np.zeros(6)
for t in range(-48, 108):
    w = sum(pmf[L] for L in range(40) if 96 <= t + 1 + L <= 107)
    if w == 0:
        continue
    d = dep[:, t + 48] if t <= 92 else future
    founders27 += (d * yield_ * w) @ share
founded = np.zeros((6, 156))
for l in links:
    founded[l["dest"], l["land"] + 48] += 1
organic27 = count_model(counts - founded)
total27 = organic27 + founders27
order = list(np.argsort(-total27))
ok("V8", order[0] == 0 and order[1] == 2, "the call: %s %.1f, then %s %.1f; organic %.1f, founded by leavers %.1f; median lag %d "
   "months; %d links, yield %.4f per removal" % (REGIONS[order[0]], total27[order[0]], REGIONS[order[1]], total27[order[1]],
                                                 organic27[0], founders27[0], int(np.median([l["lag"] for l in links if l["land"] <= 92])),
                                                 len(links), yield_))

# ------------------------------------------------------------------------------------------------- the asks
memo_in = texts["memo_2027_onboarding_pod.docx"]
ok("V9a", "West, Central and East, which serve the Southwest and Pacific, Texas & Gulf and the Midwest, and the Southeast and "
          "Northeast" in " ".join(memo_in.split()), "the request memo names the three pods and the regions each serves")
sess = csvs["onboarding_sessions_2024-01_2026-09.csv"]
assign = defaultdict(list)
for r in csvs["onboarding_team_roster_2026-10-01.csv"]:
    assign[r["name"]].append((r["pod"], day(r["pod_from"]), day(r["pod_to"]) if r["pod_to"] else None))


def pod_of(name, d):
    return next(p for p, a, b in assign[name] if a <= d and (b is None or d <= b))


group = {"Field app training", "Office hours"}
listed_on = defaultdict(set)
for s_ in sess:
    listed_on[s_["booking_id"]].add(s_["host"])
slot = {}
for s_ in sess:
    k = (s_["host"], s_["start_utc"])
    t0 = stamp(s_["start_utc"])
    u = slot.setdefault(k, dict(host=s_["host"], t0=t0, t1=t0 + dt.timedelta(minutes=int(s_["duration_min"])), typ=s_["session_type"],
                                mins=int(s_["duration_min"]), done=False, two=False))
    u["done"] |= s_["status"] == "Completed"
    u["two"] |= len(listed_on[s_["booking_id"]]) > 1
mine = defaultdict(list)
for u in slot.values():
    if u["done"]:
        mine[u["host"]].append(u)
hours = defaultdict(float)
for host, us in mine.items():
    for u in us:
        if u["t0"].year != 2025:
            continue
        # a specialist holding a session of their own while listed on another's booking did not spend that time
        if u["two"] and any(v is not u and v["t0"] < u["t1"] and u["t0"] < v["t1"] for v in us):
            continue
        hours[(pod_of(host, u["t0"].date()), (u["t0"].month - 1) // 3 + 1, "group" if u["typ"] in group else "one-to-one")] += u["mins"] / 60
cat = csvs["product_catalog_2026-07-01.csv"]
amend = defaultdict(list)
ah = amend_sheet[1]
for v in amend_sheet[2:]:
    a = dict(zip(ah, v))
    amend[a["Subscription Number"]].append(a)


def plan_by_price(seats, acv):
    """the plan whose list price for these seats, less 0, 5, 10 or 15 per cent, is this contract value"""
    hit = []
    for c in cat:
        base = (float(c["list_price_month_usd"]) + max(0, seats - int(c["seats_included"])) * float(c["extra_seat_month_usd"])) * 12
        if any(abs(acv - round(base * (1 - d), 2)) < 0.011 for d in (0, 0.05, 0.10, 0.15)):
            hit.append(c["plan"])
    return hit


def terms(a, side):
    return (a["SKU " + side], int(a["Seats " + side]), round(float(a["ACV %s (USD)" % side]), 2))


plan_n = defaultdict(int)
acv_k = defaultdict(float)
sku_seen = defaultdict(lambda: defaultdict(int))
starts = []
for e in events:
    if e["first"].year != 2025:
        continue
    s = e["sub"]
    now = (s["plan"], int(s["seats"]), round(float(s["acv"]), 2))
    later = [a for a in amend.get(s["reg"], []) if day(a["Effective Date"]) > e["first"] and terms(a, "Before") != terms(a, "After")]
    # undo the later amendments from the latest back, each day's in the order their terms either side chain
    while later:
        last_day = max(day(a["Effective Date"]) for a in later)
        a = next(a for a in later if day(a["Effective Date"]) == last_day and terms(a, "After") == now)
        now = terms(a, "Before")
        later.remove(a)
    sku_, seats, value = now
    fit = plan_by_price(seats, value)
    for p_ in fit:
        sku_seen[sku_][p_] += 1
    starts.append((e, sku_, seats, value, fit))
# each SKU names the plan most of its starting prices fit (the numbering it was sold under); a subscription starts on the
# plan its SKU names
named = {k: max(v, key=v.get) for k, v in sku_seen.items()}
for e, sku_, seats, value, fit in starts:
    plan_n[(REGIONS[e["region"]], named[sku_])] += 1
    acv_k[(REGIONS[e["region"]], named[sku_])] += value


def core_list(seats):
    c = next(c for c in cat if c["plan"] == "Core")
    return round((float(c["list_price_month_usd"]) + max(0, seats - int(c["seats_included"])) * float(c["extra_seat_month_usd"])) * 12, 2)


odd = [x for x in starts if named[x[1]] not in x[4]]
pro_sku = [k for k, v in named.items() if v == "Pro"]
ok("V9b", len(set(named.values())) == 3 and all(len(v) == 1 for k, v in sku_seen.items() if named[k] != "Pro")
   and all(x[1] in pro_sku and x[4] == ["Core"] and abs(x[3] - core_list(x[2])) < 0.005 for x in odd),
   "each 2025 starting SKU's price fits the plan it names (%s), but for %d Pro starts priced at exactly Core's list for "
   "their seats" % ({k: dict(v) for k, v in sku_seen.items()}, len(odd)))

# ------------------------------------------------------------------------------------------------- against the bundle
figs = json.load(open(os.path.join(HERE, "figures.json")))
ok("V9", figs["call"]["region"] == REGIONS[order[0]] and figs["call"]["second"] == REGIONS[order[1]]
   and figs["call"]["figure_rounded"] == int(round(total27[order[0]] / 200) * 200) and abs(figs["call"]["figure"] - total27[order[0]]) < 0.05,
   "figures.json's call %s %s (%.1f), second %s" % (figs["call"]["region"], figs["call"]["figure_rounded"], figs["call"]["figure"],
                                                    figs["call"]["second"]))
ok("V10", all(abs(figs["hours"]["%s Q%d %s" % k] - round(v, 2)) < 1e-9 for k, v in hours.items()) and len(hours) == 24,
   "all 24 pod-hour cells recomputed")
ok("V11", all(figs["cohort_counts"]["%s %s" % k] == v for k, v in plan_n.items())
   and all(abs(figs["cohort_acv_k"]["%s %s" % k] - round(v / 1000, 1)) < 1e-9 for k, v in acv_k.items()) and len(acv_k) == 18,
   "all 18 starting-plan counts and 18 signed contract values recomputed")
book = os.path.join(GOLDEN, "onboarding_staffing_2025.xlsx")
hs = sheet(book, "Specialist hours 2025")
cell_h = {(r[0], q, r[1].lower()): r[1 + q] for r in hs if r and r[0] in ("West", "Central", "East") for q in range(1, 5)}
ok("V12", all(abs(cell_h[k] - round(v, 2)) < 1e-9 for k, v in hours.items()), "the workbook's hours sheet agrees")
cs = sheet(book, "2025 new customers")
cell_c = {r[0]: r for r in cs if r and r[0] in REGIONS}
P3 = ["Core", "Pro", "Enterprise"]
ok("V13", all(cell_c[r][1 + P3.index(p)] == v for (r, p), v in plan_n.items())
   and all(abs(cell_c[r][5 + P3.index(p)] - round(v / 1000, 1)) < 1e-9 for (r, p), v in acv_k.items()),
   "the workbook's 2025 customers sheet agrees")
memo = " ".join(p.text for p in docx.Document(os.path.join(GOLDEN, "onboarding_pod_2027.docx")).paragraphs)
ok("V14", ("with the %s team" % REGIONS[order[0]]) in memo and ("about %d new paying customers" % figs["call"]["figure_rounded"]) in memo
   and ("%s comes second" % REGIONS[order[1]]) in memo, "the memo names the region, the figure and the second")
sub = os.path.join(TASK, "submission.md")
if os.path.exists(sub):
    text = open(sub, encoding="utf-8").read()
    want = ["%s" % figs["call"]["figure_rounded"]] + ["{:,.2f}".format(round(v, 2)) for v in hours.values()] + \
           ["{:,.1f}".format(round(v / 1000, 1)) for v in acv_k.values()]
    missing = [w for w in want if w not in text]
    ok("V15", not missing, "submission.md states every recomputed figure %s" % missing[:5])
print("%d FAIL" % len(FAILS))
sys.exit(1 if FAILS else 0)
