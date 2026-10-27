"""Every route on the shipped files under target/ (nothing from the generator's state): the new paying customers by
region and month under the memo's count, the careful forecasts and the ladder's rungs, the founder pipeline and the
golden family, and the asks. mess= names a Layer 1 item handled carelessly (for R2); the golden is mess=().

    python3 calc.py <target_dir>"""
from __future__ import annotations

import csv
import datetime as dt
import math
import os
import re
import sys
from collections import defaultdict

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import models as M      # noqa: E402

REG = ["Southwest", "Southeast", "Texas & Gulf", "Midwest", "Northeast", "Pacific"]
NICK = {"Bob": "Robert", "Bill": "William", "Rick": "Richard", "Joe": "Joseph", "Tom": "Thomas", "Chris": "Christopher",
        "Dan": "Daniel", "Matt": "Matthew", "Tony": "Anthony", "Mike": "Michael", "Jim": "James", "Steve": "Steven",
        "Andy": "Andrew", "Ken": "Kenneth", "Tim": "Timothy", "Ron": "Ronald", "Ed": "Edward", "Jeff": "Jeffrey",
        "Nick": "Nicholas", "Jon": "Jonathan", "Greg": "Gregory", "Alex": "Alexander", "Pat": "Patrick", "Liz": "Elizabeth",
        "Jen": "Jennifer", "Kate": "Katherine", "Becky": "Rebecca", "Sam": "Samuel", "Debbie": "Deborah", "Kim": "Kimberly",
        "Ben": "Benjamin", "Zach": "Zachary"}
PROJECT_TITLES = re.compile(r"project manager|\bpm\b|project mgr|superintendent|supt|\bsuper\b|project engineer|\bpe\b|"
                            r"field engineer|project coordinator", re.I)
F = dict(legacy="billing_legacy_subscriptions_2015-2023.csv", reg="billing_subscription_register_2022-12_2026-09.xlsx",
         crm="crm_accounts_2026-10-01.csv", dir="product_user_directory_2026-10-01.csv", terr="sales_territories_2024-03.csv",
         spend="marketing_spend_FY2020-FY2028.xlsx", cases="onboarding_cases_2024-01_2026-09.csv",
         sessions="onboarding_sessions_2024-01_2026-09.csv", roster="onboarding_team_roster_2026-10-01.csv",
         catalog="product_catalog_2026-07-01.csv")
POD_OF = {"Southwest": "West", "Pacific": "West", "Texas & Gulf": "Central", "Midwest": "Central", "Southeast": "East",
          "Northeast": "East"}
PODS = ("West", "Central", "East")
GROUP_TYPES = ("Field app training", "Office hours")
END = dt.date(2026, 9, 30)


def tindex(d):
    return (d.year - 2019) * 12 + d.month - 1


def _date(v):
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return dt.date.fromisoformat(str(v)[:10])


def _ts(v):
    return dt.datetime.strptime(v, "%Y-%m-%dT%H:%M:%SZ") if v else None


# ================================================================================================ billing and CRM
def load(T, mess=()):
    import openpyxl
    terr = {}
    for r in csv.DictReader(open(os.path.join(T, F["terr"]))):
        terr[r["state"]] = r["sales_region_before_2024_03_01"] if "old_map" in mess else r["sales_region"]
    crm = list(csv.DictReader(open(os.path.join(T, F["crm"]))))
    by_id = {r["Account ID"]: r for r in crm}

    def root(aid):
        seen = 0
        while by_id[aid]["Parent Account ID"] and seen < 5:
            aid = by_id[aid]["Parent Account ID"]
            seen += 1
        return aid
    of_code, of_acct = {}, {}
    for r in crm:
        if r["Legacy Customer Code"]:
            of_code[r["Legacy Customer Code"]] = r["Account ID"]
        if r["Billing Account Number"]:
            of_acct[r["Billing Account Number"]] = r["Account ID"]
    subs = []
    for r in csv.DictReader(open(os.path.join(T, F["legacy"]))):
        subs.append(dict(src="legacy", sid=r["subscription_id"], acct=r["customer_code"], crm=of_code.get(r["customer_code"]),
                         name=r["company"], state=r["billing_state"], first=_date(r["first_invoice_date"]),
                         paid=r["first_invoice_status"] == "paid", cancel=_date(r["cancelled_on"]), parent_bill="",
                         migrated="", plan=r["plan"], acv=int(r["mrr_cents"]) * 12 / 100.0, last_amend=None))
    wb = openpyxl.load_workbook(os.path.join(T, F["reg"]), read_only=True)
    rows = list(wb["Subscriptions"].iter_rows(values_only=True))
    hdr = rows[1]
    for r in rows[2:]:
        d = dict(zip(hdr, r))
        if not d["Account Number"]:
            continue
        subs.append(dict(src="register", sid=d["Subscription Number"], acct=d["Account Number"], crm=of_acct.get(d["Account Number"]),
                         name=d["Account Name"], state=d["Bill To State"], first=_date(d["First Invoice Date"]),
                         paid=d["First Invoice Status"] == "Paid", cancel=_date(d["Cancellation Date"]),
                         parent_bill=d["Parent Account Number"] or "", migrated=d["Migrated From (legacy subscription)"] or "",
                         plan=d["Rate Plan SKU"], seats=d["Seats"], acv=float(d["ACV (USD)"]), last_amend=_date(d["Last Amended"])))
    wb.close()
    if "both_systems" not in mess:
        leg = {s["sid"]: s for s in subs if s["src"] == "legacy"}
        for s in subs:
            if s["src"] == "register" and s["migrated"] in leg:
                L = leg[s["migrated"]]
                if s["cancel"] and not L["cancel"]:
                    L["cancel"] = s["cancel"]
                L["crm"] = L["crm"] or s["crm"]
                L["reg"] = s
        subs = [s for s in subs if not (s["src"] == "register" and s["migrated"] in leg)]
    if "voided" not in mess:
        subs = [s for s in subs if s["paid"]]
    # company = root CRM account; a subscription billed to a division rolls up to it
    comp = defaultdict(list)
    for s in subs:
        if not s["crm"]:
            continue
        rid = root(s["crm"])
        s["root"] = rid
        s["division"] = s["crm"] != rid
        comp[rid].append(s)
    events = []          # new paying customer events: (root, first invoice date, region, subscription)
    for rid, ss in comp.items():
        own = sorted([s for s in ss if not s["division"]] if "divisions" not in mess else ss, key=lambda s: (s["first"], s["sid"]))
        last_cancel, active_until = None, None
        for k, s in enumerate(own):
            new = False
            if k == 0:
                new = True
            elif "reactivations" in mess:
                new = True
            elif "returns_as_reactivations" in mess:
                new = False
            else:
                # a return more than 12 months after the company's last cancellation, nothing in force between
                lc = max((x["cancel"] for x in own[:k] if x["cancel"] is not None), default=None)
                in_force = any(x["cancel"] is None or x["cancel"] >= s["first"] for x in own[:k])
                new = (lc is not None and not in_force and (s["first"] - lc).days > 366)
            if new:
                events.append(dict(root=rid, first=s["first"], state=by_id[rid]["Billing State"], sub=s))
    new = np.zeros((6, 156), int)          # months -48 .. 107
    ev = []
    for e in events:
        if e["first"] > END or e["first"] < dt.date(2015, 1, 1):
            continue
        if "crm_region" in mess:
            reg = by_id[e["root"]]["Sales Region"]
        else:
            reg = terr.get(e["state"])
        if reg is None:
            continue
        e["region"] = REG.index(reg)
        e["t"] = tindex(e["first"])
        e["ws"] = by_id[e["root"]]["Product Workspace ID"]
        new[e["region"], e["t"] + 48] += 1
        ev.append(e)
    region_of_root = {}
    for rid in comp:
        st = by_id[rid]["Billing State"]
        if st in terr:
            region_of_root[rid] = REG.index(terr[st])
    ws_root = {}
    for r in crm:
        if r["Product Workspace ID"] and not r["Parent Account ID"]:
            ws_root[r["Product Workspace ID"]] = r["Account ID"]
    cancels = defaultdict(list)
    first_paid = {}
    for rid, ss in comp.items():
        for s in ss:
            if s["cancel"]:
                cancels[rid].append(s["cancel"])
        own = [s["first"] for s in ss if not s["division"]]
        if own:
            first_paid[rid] = min(own)
    return dict(T=T, terr=terr, crm=crm, by_id=by_id, subs=subs, comp=comp, events=ev, new=new,
                region_of_root=region_of_root, ws_root=ws_root, cancels=cancels, first_paid=first_paid, mess=mess)


def spend(T, mess=()):
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(T, F["spend"]), read_only=True)
    act = defaultdict(float)
    for r in list(wb["Actuals"].iter_rows(values_only=True))[1:]:
        fy, per, month, region, ch, v = r
        y, m = map(int, month.split("-"))
        act[(REG.index(region), (y - 2019) * 12 + m - 1)] += float(v)
    plan = {}
    for r in list(wb["Plan Oct-2026 to Jan-2028"].iter_rows(values_only=True))[1:]:
        fy, per, month, region, v = r
        if "fiscal_as_calendar" in mess:
            # the plan's periods read as calendar months: FY2028 P01 to P12 taken as January to December 2027
            y_, p_ = int(fy[2:]), int(per[1:])
            t = (y_ - 1 - 2019) * 12 + p_ - 1
        else:
            y, m = map(int, month.split("-"))
            t = (y - 2019) * 12 + m - 1
        plan[(REG.index(region), t)] = float(v)
    wb.close()
    sp = {}
    for k, v in act.items():
        sp[k] = v
    sp.update(plan)
    return [lambda t, r=r: sp[(r, t)] for r in range(6)]


# ================================================================================================ the user directory
LINK = dict(pre_paid_days=60, closure_days=3, window_days=36 * 31, pick="latest", founder="first_user", name_scope="region")


def directory(D, mess=(), **opt):
    o = dict(LINK, **opt)
    users = opt.get("_users") or list(csv.DictReader(open(os.path.join(D["T"], F["dir"]))))
    by_ws = defaultdict(list)
    for u in users:
        by_ws[u["workspace_id"]].append(u)

    def phone_key(p):
        if "phone_raw" in mess:
            return p.strip()
        d_ = re.sub(r"\D", "", p)
        return d_[-10:] if len(d_) >= 10 else ""

    def name_key(n):
        parts = [x for x in re.sub(r"[^A-Za-z ]", " ", n).split() if len(x) > 1]
        if not parts:
            return ""
        return "%s %s" % (NICK.get(parts[0], parts[0]).lower(), parts[-1].lower())
    rem = []
    for u in users:
        if not u["deactivated_at"]:
            continue
        rid = D["ws_root"].get(u["workspace_id"])
        if rid is None or rid not in D["region_of_root"]:
            continue
        da = _ts(u["deactivated_at"])
        dd = da.date() if "utc_dates" in mess else (da - dt.timedelta(hours=7)).date()
        if rid in D["first_paid"] and dd < D["first_paid"][rid] - dt.timedelta(days=o["pre_paid_days"]):
            continue
        closure = u["deactivation_reason"] == "Account closed" or any(abs((da.date() - c).days) <= o["closure_days"]
                                                                      for c in D["cancels"].get(rid, []))
        if closure and "closed_as_removal" not in mess:
            continue
        if "project_titles" in mess and not PROJECT_TITLES.search(u["job_title"] or ""):
            continue
        rem.append(dict(ws=u["workspace_id"], root=rid, region=D["region_of_root"][rid], t=tindex(dd), date=dd,
                        phone=phone_key(u["mobile"]), name=name_key(u["full_name"]), reason=u["deactivation_reason"]))
    dep = [defaultdict(int) for _ in range(6)]
    for x in rem:
        if x["t"] >= -48:
            dep[x["region"]][x["t"]] += 1
    by_phone, by_name = defaultdict(list), defaultdict(list)
    for x in rem:
        if x["phone"]:
            by_phone[x["phone"]].append(x)
        if x["name"]:
            by_name[x["name"]].append(x)
    links = []
    founder_t = np.zeros((6, 156), int)
    for e in D["events"]:
        us = sorted(by_ws.get(e["ws"], []), key=lambda u: (u["created_at"], u["user_id"]))
        if o["founder"] == "first_admin":
            us = [u for u in us if u["workspace_role"] == "Admin"] or us
        if not us:
            continue
        f = us[0]
        pk = phone_key(f["mobile"])
        cands = []
        if pk:
            cands = [x for x in by_phone.get(pk, []) if x["ws"] != e["ws"]]
        elif "phone_only" not in mess:
            nk = name_key(f["full_name"])
            cands = [x for x in by_name.get(nk, []) if x["ws"] != e["ws"] and (x["region"] == e["region"] or o["name_scope"] == "any")]
        cands = [x for x in cands if 0 < (e["first"] - x["date"]).days <= o["window_days"]]
        if cands:
            x = max(cands, key=lambda c: c["date"]) if o["pick"] == "latest" else min(cands, key=lambda c: c["date"])
            lag = max(1, (e["first"].year - x["date"].year) * 12 + e["first"].month - x["date"].month)
            links.append(dict(lag=lag, days=(e["first"] - x["date"]).days, land_t=e["t"], dep_t=x["t"], dest=e["region"],
                              origin=x["region"], reason=x["reason"]))
            founder_t[e["region"], e["t"] + 48] += 1
    return dict(users=users, by_ws=by_ws, removals=rem, dep=dep, links=links, founders=founder_t)


# ================================================================================================ forecasts
def _series(arr, r):
    return lambda t: arr[r, t + 48]


def careful_routes(D, sp):
    """Rung 0 and its siblings: count models on each region's own new paying customers."""
    y = [_series(D["new"], r) for r in range(6)]
    out = {}
    out["poisson_region_trend_spend_2019"] = M.count_forecast(y, sp, 1, 92, True, True)
    out["poisson_common_trend_spend_2019"] = M.count_forecast(y, sp, 1, 92, False, True)
    out["poisson_common_trend_2021"] = M.count_forecast(y, None, 24, 92, False, False)
    out["poisson_region_trend_2015"] = M.count_forecast(y, None, -48, 92, True, False)
    new = D["new"]
    last12 = np.array([new[r, 81 + 48:93 + 48].sum() for r in range(6)], float)
    prev12 = np.array([new[r, 69 + 48:81 + 48].sum() for r in range(6)], float)
    g = last12 / np.maximum(prev12, 1)
    out["last12_times_growth"] = last12 * g ** (15 / 12)
    return out


def _ols_trend(y, start, end, horizon=range(96, 108)):
    """Each region's monthly count on its own linear trend and month of year (least squares)."""
    out = np.zeros(6)
    for r in range(6):
        ts = list(range(start, end + 1))
        X = np.array([[1.0, t / 12.0] + [1.0 if t % 12 == m else 0.0 for m in range(1, 12)] for t in ts])
        b = np.linalg.lstsq(X, np.array([y[r](t) for t in ts], float), rcond=None)[0]
        out[r] = sum(max(0.0, float(np.dot([1.0, t / 12.0] + [1.0 if t % 12 == m else 0.0 for m in range(1, 12)], b)))
                     for t in horizon)
    return out


def probe_routes(D, sp):
    """The two call-only probes' routes, run on these files: p166ya, a four-fold back-test choosing among Poisson count
    models (region or common trend, with or without spend, fitted from 2019, 2021 or 2023) and naive rates; p166yb, an
    equal-weight ensemble of a log-linear Poisson trend and a linear trend, each with month of year."""
    y = [_series(D["new"], r) for r in range(6)]
    cands = {}
    for rt in (True, False):
        for us in (True, False):
            for st in (1, 24, 48):
                cands["poisson rt=%s spend=%s from %d" % (rt, us, st)] = (lambda e, h, rt=rt, us=us, st=st:
                                                                        M.count_forecast(y, sp, st, e, rt, us, h))
    cands["linear trend from 2019"] = lambda e, h: _ols_trend(y, 1, e, h)
    cands["linear trend from 2021"] = lambda e, h: _ols_trend(y, 24, e, h)
    err = {}
    for k, f in cands.items():
        e_ = 0.0
        for origin in (56, 68, 80, 86):                     # Sep 2023, Sep 2024, Sep 2025, Mar 2026: next six or twelve months
            h = range(origin + 1, min(origin + 13, 93))
            fc = f(origin, h)
            act = np.array([sum(y[r](t) for t in h) for r in range(6)], float)
            e_ += float(np.abs(fc - act).sum())
        err[k] = e_
    best = min(err, key=err.get)
    ya = cands[best](92, range(96, 108))
    yb = 0.5 * M.count_forecast(y, sp, 1, 92, True, True) + 0.5 * _ols_trend(y, 1, 92)
    return dict(ya=ya, ya_choice=best, ya_err=err, yb=yb)


def rung2_installed_users(D, Y, sp):
    """Word of mouth through the installed base: new customers on trend, season, spend and the log of active users on
    paying workspaces six months earlier."""
    active = np.zeros((6, 156))
    for u in Y["users"]:
        rid = D["ws_root"].get(u["workspace_id"])
        if rid is None or rid not in D["region_of_root"]:
            continue
        r = D["region_of_root"][rid]
        t0 = tindex(_ts(u["created_at"]).date())
        t1 = tindex(_ts(u["deactivated_at"]).date()) if u["deactivated_at"] else 200
        for t in range(max(t0, -48), min(t1, 93)):
            active[r, t + 48] += 1
    import math as _m
    y = [_series(D["new"], r) for r in range(6)]
    ts = list(range(1, 93))
    sbar = [float(np.mean([sp[r](t) for t in ts])) for r in range(6)]
    abar = [float(np.mean([active[r, t - 6 + 48] for t in ts])) for r in range(6)]
    rows, yy = [], []
    for r in range(6):
        for t in ts:
            x = [0.0] * 6
            x[r] = 1.0
            m = [0.0] * 11
            if t % 12 > 0:
                m[t % 12 - 1] = 1.0
            rows.append(x + [t / 12.0] + m + [_m.log(sp[r](t) / sbar[r]), _m.log(max(active[r, t - 6 + 48], 1) / abar[r])])
            yy.append(y[r](t))
    b = M.poisson_glm(np.array(yy, float), np.array(rows))
    out = np.zeros(6)
    for r in range(6):
        for t in range(96, 108):
            a = active[r, min(t - 6, 92) + 48]
            x = [0.0] * 6
            x[r] = 1.0
            m = [0.0] * 11
            if t % 12 > 0:
                m[t % 12 - 1] = 1.0
            row = np.array(x + [t / 12.0] + m + [_m.log(sp[r](t) / sbar[r]), _m.log(max(a, 1) / abar[r])])
            out[r] += math.exp(row @ b)
    return out, b[-1], active


def referral_share(D):
    """The CRM's customer-referral lead source as a share of each region's new customers by year."""
    out = {}
    for e in D["events"]:
        y = e["first"].year
        if y < 2019:
            continue
        lead = D["by_id"][e["root"]]["Lead Source"]
        k = (e["region"], y)
        n, k_ = out.get(k, (0, 0))
        out[k] = (n + 1, k_ + (lead == "Customer referral"))
    return {k: v[1] / v[0] for k, v in out.items()}


def golden_routes(D, Y, sp, lagform="emp", rateform="pooled", destform="same", organic="rt_spend", fut_form="basis"):
    """Organic count model on the customers no linked leaver founded, plus the founders from leavers: the yield per
    leaver, the lag and the destination measured on the links, realised leavers to September 2026 and the planning
    basis's leavers (each region's 2025 monthly average) after."""
    links = Y["links"]
    lags = [l["lag"] for l in links if l["land_t"] <= 92]
    if lagform == "emp":
        pmf = M.lag_pmf_from_links(lags)
    elif lagform == "gamma":
        pmf = M.gamma_pmf_from_links(lags)
    elif lagform == "emp_recent":
        pmf = M.lag_pmf_from_links([l["lag"] for l in links if 48 <= l["land_t"] <= 92])
    if destform == "same":
        same = float(np.mean([l["dest"] == l["origin"] for l in links]))
        S = np.full((6, 6), (1 - same) / 5)
        np.fill_diagonal(S, same)
    else:
        S = np.zeros((6, 6))
        for l in links:
            S[l["origin"], l["dest"]] += 1
        S = S / S.sum(1, keepdims=True)
    dep = Y["dep"]
    tot = defaultdict(int)
    for r in range(6):
        for t, v in dep[r].items():
            tot[t] += v
    t_lo = -36 if rateform != "recent" else 48
    if rateform in ("pooled", "recent"):
        land = sum(1 for l in links if t_lo <= l["land_t"] <= 92)
        p = M.founder_yield(land, tot, pmf, t_lo, 92)
    else:
        p = np.zeros(6)
        for r in range(6):
            land = sum(1 for l in links if l["origin"] == r and -36 <= l["land_t"] <= 92)
            p[r] = M.founder_yield(land, dict(dep[r]), pmf, -36, 92)
    if fut_form == "basis":
        fut = np.array([sum(dep[r].get(t, 0) for t in range(72, 84)) / 12.0 for r in range(6)])
    elif fut_form == "last12":
        fut = np.array([sum(dep[r].get(t, 0) for t in range(81, 93)) / 12.0 for r in range(6)])
    else:
        fut = np.zeros(6)
    ch = M.founder_2027(dep, fut, pmf, p, S)
    org = D["new"] - Y["founders"]
    yo = [_series(org, r) for r in range(6)]
    if organic == "rt_spend":
        o = M.count_forecast(yo, sp, 1, 92, True, True)
    elif organic == "ct_spend":
        o = M.count_forecast(yo, sp, 1, 92, False, True)
    elif organic == "rt_nospend":
        o = M.count_forecast(yo, None, 1, 92, True, False)
    elif organic == "ct_2021":
        o = M.count_forecast(yo, None, 24, 92, False, False)
    elif organic == "rt_2015":
        o = M.count_forecast(yo, None, -48, 92, True, False)
    return o + ch, dict(p=p, pmf=pmf, S=S, fut=fut, champ=ch, org=o, n_links=len(links))


FAMILY = [(lag, rate, dest, org) for org in ("rt_spend", "ct_spend", "rt_nospend", "ct_2021", "rt_2015")
          for lag in ("emp", "gamma", "emp_recent") for rate in ("pooled", "recent", "by_origin") for dest in ("same", "od")]


def name_fallback_links(D, Y, scope="region"):
    """links by mobile, and by name (nicknames folded) where the founder's mobile finds no earlier removal, within the
    founder's region or anywhere"""
    by_phone, by_name = defaultdict(list), defaultdict(list)
    for x in Y["removals"]:
        if x["phone"]:
            by_phone[x["phone"]].append(x)
        if x["name"]:
            by_name[x["name"]].append(x)

    def name_key(n):
        parts = [p_ for p_ in re.sub(r"[^A-Za-z ]", " ", n).split() if len(p_) > 1]
        return "%s %s" % (NICK.get(parts[0], parts[0]).lower(), parts[-1].lower()) if parts else ""
    links, founders = [], np.zeros((6, 156), int)
    for e in D["events"]:
        us = sorted(Y["by_ws"].get(e["ws"], []), key=lambda u: (u["created_at"], u["user_id"]))
        if not us:
            continue
        f = us[0]
        d_ = re.sub(r"\D", "", f["mobile"])
        pk = d_[-10:] if len(d_) >= 10 else ""
        win = LINK["window_days"]
        cands = [x for x in by_phone.get(pk, []) if x["ws"] != e["ws"] and 0 < (e["first"] - x["date"]).days <= win]
        if not cands:
            cands = [x for x in by_name.get(name_key(f["full_name"]), []) if x["ws"] != e["ws"]
                     and (scope == "any" or x["region"] == e["region"]) and 0 < (e["first"] - x["date"]).days <= win]
        if cands:
            x = max(cands, key=lambda c: c["date"])
            lag = max(1, (e["first"].year - x["date"].year) * 12 + e["first"].month - x["date"].month)
            links.append(dict(lag=lag, land_t=e["t"], dep_t=x["t"], dest=e["region"], origin=x["region"]))
            founders[e["region"], e["t"] + 48] += 1
    return dict(Y, links=links, founders=founders)


def judge_route(D, Yv, sp, link_from=-48, org_start=1, rate="pooled", lagshift=0, parts=False):
    """one of the round 1 judge's constructions: founders sought only from link_from (the organic series cleaned only
    there), the organic fit from org_start, a pooled yield or a mature-cohort yield (founders per removal of January 2023
    to October 2024), the linked lags shifted by lagshift months; parts= also returns the organic and founder-led parts"""
    links = [l_ for l_ in Yv["links"] if l_["land_t"] >= link_from]
    pmf = M.lag_pmf_from_links([max(1, l_["lag"] + lagshift) for l_ in links if l_["land_t"] <= 92])
    tot = defaultdict(int)
    for r in range(6):
        for t, v in Yv["dep"][r].items():
            tot[t] += v
    lo = max(-36, link_from)
    if rate == "pooled":
        p = M.founder_yield(sum(1 for l_ in links if lo <= l_["land_t"] <= 92), tot, pmf, lo, 92)
    else:
        p = sum(1 for l_ in links if 48 <= l_["dep_t"] <= 69) / sum(v for t, v in tot.items() if 48 <= t <= 69)
    same = float(np.mean([l_["dest"] == l_["origin"] for l_ in links]))
    S = np.full((6, 6), (1 - same) / 5)
    np.fill_diagonal(S, same)
    fut = np.array([sum(Yv["dep"][r].get(t, 0) for t in range(72, 84)) / 12.0 for r in range(6)])
    ch = M.founder_2027(Yv["dep"], fut, pmf, p, S)
    fz = np.zeros_like(Yv["founders"])
    for l_ in links:
        fz[l_["dest"], l_["land_t"] + 48] += 1
    yo = [_series(D["new"] - fz, r) for r in range(6)]
    o = M.count_forecast(yo, sp, org_start, 92, True, True)
    return (o + ch, o, ch) if parts else o + ch


def judge_family(D, Y, sp):
    """the round 1 judge's 72 choices: three linkers (mobile; mobile, then a name within the region; mobile, then a name
    anywhere) x founders sought from 2015 or only from 2023 x the organic fit from 2019, 2021 or 2023 x a pooled or a
    mature yield x the lag as linked or a month shorter"""
    out = {}
    for vn, Yv in (("mobile", Y), ("name_region", name_fallback_links(D, Y, "region")),
                   ("name_any", name_fallback_links(D, Y, "any"))):
        for link_from in (-48, 48):
            for org_start in (1, 24, 48):
                for rate in ("pooled", "mature"):
                    for ls in (0, -1):
                        out[(vn, link_from, org_start, rate, ls)] = judge_route(D, Yv, sp, link_from, org_start, rate, ls)
    return out


def golden(T, mess=()):
    D = load(T, mess)
    sp = spend(T, mess)
    Y = directory(D, mess)
    g, info = golden_routes(D, Y, sp)
    return D, sp, Y, g, info


# ================================================================================================ the asks
def _region_of_crm(D):
    out = {}
    for r in D["crm"]:
        st = r["Billing State"]
        out[r["Account ID"]] = REG.index(D["terr"][st]) if st in D["terr"] else None
    return out


def roster_history(T):
    """each specialist's pod assignments from the roster: name -> [(pod, first day, last day or None)]"""
    out = defaultdict(list)
    for r in csv.DictReader(open(os.path.join(T, F["roster"]))):
        out[r["name"]].append((r["pod"], dt.date.fromisoformat(r["pod_from"]),
                               dt.date.fromisoformat(r["pod_to"]) if r["pod_to"] else None))
    return dict(out)


def pod_on(hist, name, d):
    for pod, a, b in hist[name]:
        if a <= d and (b is None or d <= b):
            return pod
    return None


def _held_sessions(T, D):
    """sessions as the scheduling tool lists them: one per host and start, held when any of its bookings was completed,
    with the customer's region by the case and by the booking's account, whether the booking lists a second specialist,
    and whether the host held another session overlapping it"""
    case_acct = {r["case_id"]: r["crm_account_id"] for r in csv.DictReader(open(os.path.join(T, F["cases"])))}
    rows = list(csv.DictReader(open(os.path.join(T, F["sessions"]))))
    reg_of = _region_of_crm(D)
    hosts_of = defaultdict(set)
    for s in rows:
        hosts_of[s["booking_id"]].add(s["host"])
    held = {(s["host"], s["start_utc"]) for s in rows if s["status"] == "Completed"}
    out, seen = [], {}
    for s in rows:
        k = (s["host"], s["start_utc"])
        if k in seen:
            seen[k]["shared"] |= len(hosts_of[s["booking_id"]]) > 1
            continue
        st = _ts(s["start_utc"])
        seen[k] = dict(host=s["host"], start=st, end=st + dt.timedelta(minutes=int(s["duration_min"])), month=(st.year, st.month),
                       q=(st.month - 1) // 3 + 1, year=st.year, held=k in held, hours=int(s["duration_min"]) / 60.0,
                       kind="group" if s["session_type"] in GROUP_TYPES else "one-to-one", bk=s["booking_id"],
                       case_reg=reg_of.get(case_acct.get(s["case_id"])), booking_reg=reg_of.get(s["crm_account_id"]),
                       shared=len(hosts_of[s["booking_id"]]) > 1)
        out.append(seen[k])
    by_host = defaultdict(list)
    for x in out:
        if x["held"]:
            by_host[x["host"]].append(x)
    for xs in by_host.values():
        xs.sort(key=lambda x: x["start"])
        for i, x in enumerate(xs):
            x["clash"] = any(y is not x and y["start"] < x["end"] and x["start"] < y["end"]
                             for y in xs[max(0, i - 4):i + 5])
    for x in out:
        x.setdefault("clash", False)
    return rows, out


HOURS_MODES = ("golden", "listed", "minutes", "roster_now", "customer", "booking_account", "booking_dedup")


def ask_pod_hours(T, mode="golden", D=None):
    """Hours of specialist time each pod's specialists spent in one-to-one and in group sessions, by quarter of 2025.

    golden: one session per host and start, held when any booking completed, its length once for that host, in the pod the
    roster has the host in on the session's date; a specialist listed on a booking that another specialist also hosts,
    while holding a session of their own at that time, did not spend that time (the booking is the other's).
    listed: every listed host's session counted, the roster's pods by date. minutes: each specialist's held minutes once,
    the roster's pods by date (equal to the golden: every listed-not-hosting booking lies wholly inside the specialist's
    own session). roster_now: the same as listed, each host in the pod the
    roster has them in today (or when they left). customer: each session in the pod of its case's customer.
    booking_account: in the pod of the account the booking was made under. booking_dedup: one host per booking reference
    (the first listed in the export), the roster's pods by date."""
    D = D or load(T)
    rows, sess = _held_sessions(T, D)
    hist = roster_history(T)
    now = {n: sorted(v, key=lambda x: x[1])[-1][0] for n, v in hist.items()}
    out = defaultdict(float)
    if mode == "booking_dedup":
        first = {}
        for s in rows:
            first.setdefault(s["booking_id"], s["host"])
    if mode == "minutes":
        # each specialist's held time counted once: a session's minutes already covered by a longer session of theirs drop
        by_host = defaultdict(list)
        for x in sess:
            if x["held"]:
                by_host[x["host"]].append(x)
        for xs in by_host.values():
            done = []
            for x in sorted(xs, key=lambda x: (-(x["end"] - x["start"]), x["start"])):
                cov = sum(max(0.0, (min(x["end"], y["end"]) - max(x["start"], y["start"])).total_seconds()) for y in done)
                x["own_hours"] = x["hours"] - cov / 3600.0
                done.append(x)
    for x in sess:
        if x["year"] != 2025 or not x["held"]:
            continue
        if mode == "golden" and x["shared"] and x["clash"]:
            continue
        if mode == "booking_dedup" and first[x["bk"]] != x["host"]:
            continue
        if mode in ("golden", "listed", "booking_dedup", "minutes"):
            pod = pod_on(hist, x["host"], x["start"].date())
        elif mode == "roster_now":
            pod = now[x["host"]]
        elif mode == "customer":
            pod = POD_OF[REG[x["case_reg"]]] if x["case_reg"] is not None else None
        elif mode == "booking_account":
            pod = POD_OF[REG[x["booking_reg"]]] if x["booking_reg"] is not None else None
        if pod is None:
            continue
        out[(pod, x["q"], x["kind"])] += x["own_hours"] if mode == "minutes" else x["hours"]
    return dict(out)


def cohort_2025(D):
    return [e for e in D["events"] if e["first"].year == 2025]


def amendments(T):
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(T, F["reg"]), read_only=True)
    rows = list(wb["Amendments"].iter_rows(values_only=True))
    wb.close()
    hdr = rows[1]
    out = defaultdict(list)
    for r in rows[2:]:
        d = dict(zip(hdr, r))
        if d["Subscription Number"]:
            out[d["Subscription Number"]].append(dict(date=_date(d["Effective Date"]), no=d["Amendment Number"], change=d["Change"],
                                                      sku0=d["SKU Before"], sku1=d["SKU After"], seats0=d["Seats Before"],
                                                      seats1=d["Seats After"], acv0=float(d["ACV Before (USD)"]),
                                                      acv1=float(d["ACV After (USD)"])))
    for v in out.values():
        v.sort(key=lambda a: (a["date"], a["no"]))
    return out


def catalog_map(T):
    return {r["sku"]: r["plan"] for r in csv.DictReader(open(os.path.join(T, F["catalog"])))}


def _reg_row(e):
    return e["sub"]["reg"] if "reg" in e["sub"] else e["sub"]


def start_terms(e, A, on_or_after=False):
    """a cohort subscription's SKU, seats and ACV in force on its first paid invoice date: the register's terms today taken
    back through every amendment effective after that date, latest first, a day's amendments in the order their own terms
    either side chain (on_or_after: those effective on the date itself too)"""
    s = _reg_row(e)
    cur = (s["plan"], int(s["seats"]), round(float(s["acv"]), 2))
    later = [a for a in A.get(s["sid"], []) if (a["date"] >= e["first"] if on_or_after else a["date"] > e["first"])]
    for d in sorted({a["date"] for a in later}, reverse=True):
        grp = [a for a in later if a["date"] == d and (a["sku0"], a["seats0"], a["acv0"]) != (a["sku1"], a["seats1"], a["acv1"])]
        while grp:
            nxt = [a for a in grp if (a["sku1"], int(a["seats1"]), round(a["acv1"], 2)) == cur]
            if len(nxt) != 1:
                raise ValueError("amendments of %s on %s do not chain" % (s["sid"], d))
            cur = (nxt[0]["sku0"], int(nxt[0]["seats0"]), round(nxt[0]["acv0"], 2))
            grp.remove(nxt[0])
    return cur


def start_terms_by_number(e, A):
    """the terms in force on the first paid invoice date read from the amendment log in its own order: the after terms of
    the last amendment effective on or before the date, or the before terms of the first later one by date and reference"""
    s = _reg_row(e)
    am = sorted(A.get(s["sid"], []), key=lambda a: (a["date"], a["no"]))
    later = [a for a in am if a["date"] > e["first"]]
    if later:
        return later[0]["sku0"], int(later[0]["seats0"]), round(later[0]["acv0"], 2)
    return s["plan"], int(s["seats"]), round(float(s["acv"]), 2)


def price_plans(T):
    """(seats, ACV) -> the plans whose list price for those seats, less 0, 5, 10 or 15 per cent, is that contract value"""
    cat = list(csv.DictReader(open(os.path.join(T, F["catalog"]))))

    def fit(seats, acv):
        out = []
        for c in cat:
            lst = (float(c["list_price_month_usd"]) + max(0, int(seats) - int(c["seats_included"])) * float(c["extra_seat_month_usd"])) * 12
            if any(abs(acv - round(lst * (1 - d), 2)) < 0.011 for d in (0.0, 0.05, 0.10, 0.15)):
                out.append(c["plan"])
        return out
    return fit


def launch_order_map(D, A, T):
    """the SKUs as numbered before the July 2026 catalog, read off each 2025 starting subscription's own price: the plan
    whose list price for the subscription's seats, less a discount of 0, 5, 10 or 15 per cent, is its contract value"""
    fit = price_plans(T)
    fits = defaultdict(lambda: defaultdict(int))
    for e in cohort_2025(D):
        sku_, seats, acv = start_terms(e, A)
        for p_ in fit(seats, acv):
            fits[sku_][p_] += 1
    return {k: max(v, key=v.get) for k, v in fits.items()}, fits


COHORT_MODES = ("golden", "chain_price", "chain_catalog", "number_price", "number_catalog", "register_catalog", "case_catalog",
                "register_launch", "rollback_all_launch")


def ask_cohort(T, mode="golden", D=None, A=None):
    """The 2025 new paying customers by region and the plan they started on, and the annual contract value they signed
    on it, on the terms in force on the first paid invoice date.

    golden: the register's terms today taken back through every amendment effective after the first paid invoice date, a
    day's amendments in the order their terms chain, the plan the subscription's SKU of record names in the numbering it
    was sold under (launch order, read off the starting subscriptions' own prices and seats). chain_price: the same terms,
    the plan from the price and seats (reader B's route at round 3). chain_catalog: the same terms, the SKU named by today's
    catalog. number_price: the amendment log read in its own order (the
    before terms of the first later amendment by date and reference), the plan from the price and seats. number_catalog:
    the same terms, the SKU named by today's catalog. register_catalog: the register's SKU and value today, named by
    today's catalog. case_catalog: the onboarding case's SKU and value, named by today's catalog. register_launch: the
    register's terms today in launch-order names. rollback_all_launch: the amendments effective on the first invoice date
    taken back too."""
    D = D or load(T)
    A = A if A is not None else amendments(T)
    cat = catalog_map(T)
    launch, _ = launch_order_map(D, A, T)
    fit = price_plans(T)
    coh = cohort_2025(D)
    cases = defaultdict(list)
    if mode == "case_catalog":
        for r in csv.DictReader(open(os.path.join(T, F["cases"]))):
            cases[r["crm_account_id"]].append(r)
    counts = defaultdict(int)
    acv = defaultdict(float)
    for e in coh:
        s = _reg_row(e)
        if mode == "golden":
            sku_, _, value = start_terms(e, A)
            plan = launch[sku_]
        elif mode in ("chain_price", "chain_catalog"):
            sku_, seats, value = start_terms(e, A)
            plan = (fit(seats, value) or ["?"])[0] if mode == "chain_price" else cat[sku_]
        elif mode in ("number_price", "number_catalog"):
            sku_, seats, value = start_terms_by_number(e, A)
            plan = (fit(seats, value) or ["?"])[0] if mode == "number_price" else cat[sku_]
        elif mode == "rollback_all_launch":
            sku_, _, value = start_terms(e, A, on_or_after=True)
            plan = launch[sku_]
        elif mode == "register_catalog":
            plan, value = cat[s["plan"]], s["acv"]
        elif mode == "register_launch":
            plan, value = launch[s["plan"]], s["acv"]
        elif mode == "case_catalog":
            cs = sorted([c for c in cases.get(e["root"], []) if abs((dt.date.fromisoformat(c["opened_on"]) - e["first"]).days) <= 7],
                        key=lambda c: c["opened_on"])
            c = cs[0]
            plan, value = cat[c["rate_plan_sku"]], float(c["contract_acv_usd"])
        counts[(REG[e["region"]], plan)] += 1
        acv[(REG[e["region"]], plan)] += value
    return dict(counts=dict(counts), acv=dict(acv), n=len(coh))


if __name__ == "__main__":
    T = sys.argv[1]
    D, sp, Y, g, info = golden(T)
    print("new 2025 by region", D["new"][:, 72 + 48:84 + 48].sum(1).tolist())
    print("links", len(Y["links"]), "removals", len(Y["removals"]), "p", round(float(np.mean(info["p"])), 4))
    c = careful_routes(D, sp)
    for k, v in c.items():
        print("%-34s %s" % (k, ", ".join("%s %.0f" % (REG[j], v[j]) for j in np.argsort(-v)[:3])))
    print("%-34s %s" % ("golden", ", ".join("%s %.1f" % (REG[j], g[j]) for j in np.argsort(-g)[:3])))
