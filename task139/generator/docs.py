"""The pack's documents and workbooks: the signed programme charter and its September draft, the AY2027 pick memo, the
Kellerman Pryce review and its verification workbook, the triage screen's pilot register, the data dictionary, the
2027 staffing plan and the export manifest."""
from __future__ import annotations

import datetime as dt

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import LETTER                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import inch                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import common as K   # noqa: E402
import records as R  # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/"
for _n, _f in (("Carlito", "crosextra/Carlito-Regular.ttf"), ("Carlito-Bold", "crosextra/Carlito-Bold.ttf"),
               ("Carlito-Italic", "crosextra/Carlito-Italic.ttf"),
               ("Caladea", "crosextra/Caladea-Regular.ttf"), ("Caladea-Bold", "crosextra/Caladea-Bold.ttf")):
    pdfmetrics.registerFont(TTFont(_n, FONT_DIR + _f))
SANS, SANS_B, SANS_I = "Carlito", "Carlito-Bold", "Carlito-Italic"
SERIF, SERIF_B = "Caladea", "Caladea-Bold"
INK = colors.HexColor("#1d2330")
GREY = colors.HexColor("#5a6170")
NAVY = colors.HexColor("#1f3a5f")
TEAL = colors.HexColor("#2d6a6a")
RULE = colors.HexColor("#a3abb8")

BODY = ParagraphStyle("body", fontName=SERIF, fontSize=10.3, leading=13.6, textColor=INK, spaceAfter=5)
CLAUSE = ParagraphStyle("clause", parent=BODY, leftIndent=26, firstLineIndent=-26)
H1 = ParagraphStyle("h1", fontName=SANS_B, fontSize=15, leading=19, textColor=NAVY, spaceAfter=2)
H1S = ParagraphStyle("h1s", fontName=SANS, fontSize=10.0, leading=13.0, textColor=GREY, spaceAfter=8)
H2 = ParagraphStyle("h2", fontName=SANS_B, fontSize=11.0, leading=14, textColor=INK, spaceBefore=7, spaceAfter=3)
NOTE = ParagraphStyle("note", fontName=SANS_I, fontSize=8.8, leading=11.2, textColor=GREY, spaceBefore=4)
MEMO = ParagraphStyle("memo", fontName=SANS, fontSize=10.2, leading=13.4, textColor=INK, spaceAfter=1)
CELL = ParagraphStyle("cell", fontName=SANS, fontSize=9.0, leading=11.0, textColor=INK)
CELLB = ParagraphStyle("cellb", fontName=SANS_B, fontSize=9.0, leading=11.0, textColor=INK)


def _footer(left, right):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont(SANS, 7.8)
        canvas.setFillColor(GREY)
        canvas.drawString(0.85 * inch, 0.55 * inch, left)
        canvas.drawRightString(LETTER[0] - 0.85 * inch, 0.55 * inch, "%s  |  Page %d" % (right, doc.page))
        canvas.restoreState()
    return draw


def _doc(path, title, author):
    return SimpleDocTemplate(path, pagesize=LETTER, leftMargin=0.95 * inch, rightMargin=0.95 * inch,
                             topMargin=0.8 * inch, bottomMargin=0.9 * inch, title=title, author=author, subject="",
                             creator="", producer="")


def _grid(rows, widths, head=True):
    t = Table(rows, colWidths=widths, hAlign="LEFT")
    st = [("FONT", (0, 0), (-1, -1), SANS, 9), ("TEXTCOLOR", (0, 0), (-1, -1), INK),
          ("LINEBELOW", (0, 0), (-1, 0), 0.6, RULE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
          ("BOTTOMPADDING", (0, 0), (-1, -1), 3), ("TOPPADDING", (0, 0), (-1, -1), 2)]
    if head:
        st.append(("FONT", (0, 0), (-1, 0), SANS_B, 9))
    t.setStyle(TableStyle(st))
    return t


def _sec(s, title, clauses):
    s.append(Paragraph(title, H2))
    for no, text in clauses:
        s.append(Paragraph("%s&nbsp;&nbsp;&nbsp;%s" % (no, text), CLAUSE))


def _long(d):
    return "%d %s %d" % (d.day, d.strftime("%B"), d.year)


# ------------------------------------------------------------------------------------------------- the charter

def write_charter(path):
    s = [Paragraph("%s &middot; Medical Case Management" % K.COMPANY, H1S),
         Paragraph("Early Return Nursing: programme charter for 2027", H1),
         Paragraph("Approved by the Executive Committee on %s. This charter supersedes the drafts circulated in "
                   "September 2026." % _long(K.CHARTER_SIGNED), H1S)]
    _sec(s, "1 Purpose", [
        ("1.1", "Early Return Nursing (ERN) puts a nurse case manager on a lost-time claim from the day it reaches "
                "triage. The nurse calls the worker, the employer and the treating provider within a working day, "
                "coordinates care, arranges suitable duties where the employer can offer them, and stays with the "
                "file until the worker is back at full duty.")])
    _sec(s, "2 What is funded", [
        ("2.1", "The Executive Committee funds seventeen ERN teams for calendar year 2027. Each team is a nurse case "
                "manager and a return-to-work coordinator."),
        ("2.2", "A team holds twenty places: at most twenty open files at any one time.")])
    _sec(s, "3 Where the teams work", [
        ("3.1", "Each team works the intakes of one segment and sits in that segment's claim unit. Each segment's "
                "claims are worked in one unit (Appendix A)."),
        ("3.2", "A claim unit houses at most three teams."),
        ("3.3", "A unit is given no more teams than its 2027 intakes will keep in use: the place-days its intakes "
                "would hold over the year must fill every place of every team it is given."),
        ("3.4", "The Director of Medical Case Management places the teams, and the placement goes to the Committee "
                "with the 2027 plan in December.")])
    _sec(s, "4 How a team takes intakes", [
        ("4.1", "An intake is a claim whose first report of injury records time lost from work."),
        ("4.2", "Each working morning the triage screen offers a unit's new intakes to its teams in report order. A "
                "team takes an intake while it has a free place. An intake not taken by the end of its second working "
                "day stays with its adjuster.")])
    _sec(s, "5 Files and places", [
        ("5.1", "A team holds an intake's file from the day it takes it until the worker is back at full duty, the "
                "file goes back to the adjuster, or the file closes at maximum medical improvement."),
        ("5.2", "A file reopened after it closed holds a place again until it closes."),
        ("5.3", "A place-day is one place held for one day. The day a file opens and the day it closes both count.")])
    _sec(s, "6 Start and reporting", [
        ("6.1", "The teams start on Monday 4 January 2027."),
        ("6.2", "The Director reports to the Committee each quarter on places in use, files opened and closed, and "
                "closures by reason.")])
    rows = [["Segment", "Code", "Claim unit", "Office"]]
    stl, kc = "St. Louis (STL)", "Kansas City (KC)"
    office = {"HC-SS": stl, "HC-OT": stl, "RW-SS": kc, "RW-OT": kc, "MF-SS": stl, "MF-OT": kc, "CN-SS": kc, "CN-OT": stl}
    for seg in K.PROG:
        rows.append([K.NAME[seg], seg, K.UNIT[seg], office[seg]])
    s.append(KeepTogether([Paragraph("Appendix A. Claim units", H2),
                           _grid(rows, [3.2 * inch, 0.6 * inch, 1.6 * inch, 1.15 * inch])]))
    s.append(Spacer(1, 14))
    sig = [["%s" % K.PEOPLE["ceo"], "%s" % K.PEOPLE["director"]],
           ["President and Chief Executive Officer", "Director, Medical Case Management"],
           [_long(K.CHARTER_SIGNED), _long(K.CHARTER_SIGNED)]]
    t = Table(sig, colWidths=[3.2 * inch, 3.2 * inch], hAlign="LEFT")
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), SANS, 9.4), ("FONT", (0, 0), (-1, 0), SANS_B, 9.6),
                           ("LINEABOVE", (0, 0), (-1, 0), 0.6, RULE), ("TEXTCOLOR", (0, 0), (-1, -1), INK)]))
    s.append(t)
    doc = _doc(path, "Early Return Nursing programme charter 2027", K.PEOPLE["director"])
    foot = _footer("%s  |  ERN programme charter 2027 (signed)" % K.COMPANY, "MCM-2027-01")
    doc.build(s, onFirstPage=foot, onLaterPages=foot)


def write_charter_draft(path):
    import docx
    from docx.shared import Pt
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    d.add_paragraph("DRAFT v3 for discussion, %s. Not approved." % _long(K.CHARTER_DRAFT))
    d.add_heading("Early Return Nursing: programme charter for 2027", level=1)
    paras = [
        ("1 Purpose", "Early Return Nursing (ERN) puts a nurse on a lost-time claim from the day it reaches triage, to "
                      "coordinate care, arrange suitable duties where the employer can offer them and see the worker back "
                      "to full duty."),
        ("2 Funding", "The 2027 plan proposes sixteen to eighteen ERN teams, subject to the Executive Committee's "
                      "approval of the plan in October. Each team would be two nurse case managers with twenty open "
                      "files at most between them."),
        ("3 Placement", "Teams would sit in the claim units, at most three to a unit, and work that unit's intakes. "
                        "[JW: placement basis to settle with Pricing once the Kellerman Pryce review is in.]"),
        ("4 Intake", "The triage screen would offer each unit's intakes to its teams each working morning."),
        ("5 Files", "A team would hold a file until the worker is back at full duty or reaches maximum medical "
                    "improvement."),
        ("6 Start", "Teams would start on Monday 1 February 2027, after recruiting in the fourth quarter."),
    ]
    for h, t in paras:
        d.add_heading(h, level=2)
        d.add_paragraph(t)
    d.core_properties.author = K.PEOPLE["director"]
    d.core_properties.title = "ERN charter 2027 draft v3"
    d.save(path)


# ------------------------------------------------------------------------------------------------- the pick memo

def write_pick_memo(path, pick_total):
    s = [Paragraph("%s &middot; Pricing and Reserving" % K.COMPANY, H1S),
         Paragraph("Memorandum", H1),
         Paragraph("<b>To:</b> Executive Committee", MEMO),
         Paragraph("<b>From:</b> %s, Chief Pricing Actuary" % K.PEOPLE["pricing"], MEMO),
         Paragraph("<b>Copy:</b> %s, Director, Medical Case Management" % K.PEOPLE["director"], MEMO),
         Paragraph("<b>Date:</b> %s" % _long(K.MEMO_DATE), MEMO),
         Paragraph("<b>Subject:</b> AY2027 indemnity pick: the Early Return Nursing credit", MEMO),
         Spacer(1, 8)]
    paras = [
        "1. The AY2027 plan carries an indemnity pick for each segment. For the eight segments the 2027 Early Return "
        "Nursing teams can work, the pick before any credit is $%.1f million, at the pick's trend and benefit levels "
        "(plan Exhibit 3)." % (pick_total / 1e6),
        "2. The pick books the saving from the 2027 teams as one credit against these eight segments. The credit is the "
        "ultimate indemnity that the accident year 2027 intakes the teams take will not pay because the teams took "
        "them: what those intakes would pay with their adjusters alone, less what they pay with a team.",
        "3. The credit is set at run rate: every place of every team in use on every day of 2027, twenty places for 365 "
        "days. The weeks while the teams fill in January and the files still open at the end of December are not "
        "adjusted for.",
        "4. Amounts are those the pilot verified: Kellerman Pryce's verified ultimate indemnity at 30 June 2026 for "
        "temporary total, temporary partial and permanent partial disability together, open and closed claims alike, "
        "with no trend or benefit-level adjustment. The pick's severity trend and benefit-level factors apply to the "
        "whole pick, credit included, in Exhibit 4.",
        "5. Each segment's teams are credited with what the pilot measured in that segment. A segment without teams "
        "carries no credit.",
        "6. Medical is in the medical pick; this credit is indemnity only.",
        "7. I will book the credit with the placement %s supplies for the December meeting." % K.PEOPLE["director"],
    ]
    for p in paras:
        s.append(Paragraph(p, BODY))
    s.append(Spacer(1, 8))
    s.append(Paragraph("KR", BODY))
    doc = _doc(path, "AY2027 indemnity pick: the ERN credit", K.PEOPLE["pricing"])
    foot = _footer("%s  |  Pricing and Reserving" % K.COMPANY, "AY2027 pick memo 4")
    doc.build(s, onFirstPage=foot, onLaterPages=foot)


# ------------------------------------------------------------------------------------------------- the review

def write_review(path, frame_n, sample_n, open_n, total):
    s = [Paragraph("KELLERMAN PRYCE ACTUARIAL", ParagraphStyle("lh", fontName=SANS_B, fontSize=12.5, leading=15,
                                                              textColor=TEAL)),
         Paragraph("Consulting actuaries &middot; St. Louis, Missouri", H1S),
         Paragraph("Early Return Nursing pilot: verification of indemnity on pilot claims at 30 June 2026", H1),
         Paragraph("Prepared for %s &middot; %s" % (K.COMPANY, _long(K.REVIEW_DATE)), H1S)]
    _sec(s, "1 Scope", [
        ("1.1", "Haverline asked us to verify the ultimate indemnity on a sample of the claims from its Early Return "
                "Nursing (ERN) pilot, so that the amounts Haverline uses in its own evaluation of the pilot are verified "
                "figures."),
        ("1.2", "We were not asked to estimate the effect of ERN, and this report does not do so.")])
    _sec(s, "2 The pilot", [
        ("2.1", "The pilot ran on intakes triaged from 1 March 2023 to 28 February 2025 in eight claim segments. "
                "Haverline's triage screen drew intakes into the pilot and assigned each to ERN or to usual handling. "
                "Haverline's assignment register records every pilot intake and its group; we relied on it for group "
                "and segment.")])
    _sec(s, "3 Frame and sample", [
        ("3.1", "The frame was every pilot claim that had paid or reserved indemnity (temporary total, temporary partial "
                "or permanent partial disability) at 30 June 2026: %s claims." % format(frame_n, ","),
                ),
        ("3.2", "We treated each segment and assignment group as a stratum and drew %d claims from each of the sixteen "
                "strata at random, without replacement: %s claims in all." % (sample_n // 16, format(sample_n, ","))),
        ("3.3", "The workbook lists each segment's sampled claims together, in claim-number order, with our reference, "
                "the claim's status at 30 June 2026 and its verified ultimate indemnity by benefit.")])
    _sec(s, "4 Verification", [
        ("4.1", "Each sampled file was verified by a reviewer who did not know its group."),
        ("4.2", "For temporary total and temporary partial disability we checked the average weekly wage against the "
                "wage statement, the weekly rate against the Missouri maximum for the date of injury, and the periods "
                "paid against the restrictions and return-to-work records, including the three-day waiting period and "
                "its payment after the fourteenth day."),
        ("4.3", "For permanent partial disability we checked the rating, the body part and the weeks against the award "
                "or the approved settlement."),
        ("4.4", "Claims open at 30 June 2026 (%d of the sampled claims) were projected to ultimate on their case facts "
                "and our Missouri indemnity development factors." % open_n)])
    _sec(s, "5 Results", [
        ("5.1", "Verified ultimate indemnity on the %s sampled claims totals $%s. The workbook's sheet totals add to "
                "it." % (format(sample_n, ","), format(round(total, 2), ",.2f")))])
    _sec(s, "6 Limitations", [
        ("6.1", "Verified amounts are at 30 June 2026 values. We have not adjusted them for trend or for benefit changes "
                "after the dates of injury."),
        ("6.2", "Our projections for open claims are estimates; actual payments will differ.")])
    s.append(Spacer(1, 12))
    s.append(Paragraph("%s, FCAS, MAAA" % K.PEOPLE["kpa"], BODY))
    s.append(Paragraph("Principal, %s" % K.KPA, BODY))
    doc = _doc(path, "ERN pilot: verification of indemnity at 30 June 2026", K.PEOPLE["kpa"])
    foot = _footer("%s  |  Haverline ERN pilot verification" % K.KPA, "KPA-26-118")
    doc.build(s, onFirstPage=foot, onLaterPages=foot)


# ------------------------------------------------------------------------------------------------- workbooks

def _title_rows(ws, lines):
    for k, (text, style) in enumerate(lines):
        ws.append([text])
        cell = ws.cell(row=ws.max_row, column=1)
        if style == "title":
            cell.font = Font(name="Calibri", bold=True, size=12)
        elif style == "sub":
            cell.font = Font(name="Calibri", italic=True, size=10, color="595959")
        else:
            cell.font = Font(name="Calibri", size=9, color="595959")
    ws.append([])


def _head(ws, cols, widths=None, fill="D9DEE7"):
    ws.append(cols)
    row = ws.max_row
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name="Calibri", bold=True)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    return row


def write_register(path, W):
    rows = R.register_rows(W)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pilot intakes"
    _title_rows(ws, [("Early Return Nursing pilot: assignment register", "title"),
                     ("%s, Claims Operations. Triage dates 1 March 2023 to 28 February 2025." % K.COMPANY, "sub"),
                     ("One row per pilot intake, in triage order. Segment: the legacy segment code at triage.", "")])
    hr = _head(ws, ["Pilot ref", "Claim number", "Office", "Segment", "Reported", "Triaged", "Group"],
               [11, 14, 8, 9, 11, 11, 15])
    for r in rows:
        ws.append(list(r))
        ws.cell(row=ws.max_row, column=5).number_format = "yyyy-mm-dd"
        ws.cell(row=ws.max_row, column=6).number_format = "yyyy-mm-dd"
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    ws.auto_filter.ref = "A%d:G%d" % (hr, ws.max_row)
    ws2 = wb.create_sheet("Notes")
    notes = [
        "Each working morning the triage screen drew one in three of the day's intakes in the eight pilot segments "
        "into the pilot, at random.",
        "Within each segment the drawn intakes were taken in pairs, in triage order, and one of each pair was assigned at "
        "random to ERN and the other to usual handling. A segment's last drawn intake, if unpaired, was assigned at "
        "random.",
        "ERN intakes went to an ERN team on the day of triage. Usual-handling intakes stayed with their adjusters.",
        "An intake is a claim whose first report of injury records time lost from work.",
        "Segment codes are the legacy codes in use at triage; see the 2025 segment code crosswalk.",
        "Prepared by %s, Claims Operations, 7 March 2025." % K.PEOPLE["claimsops"],
    ]
    ws2.append(["Notes"])
    ws2["A1"].font = Font(name="Calibri", bold=True, size=12)
    for n in notes:
        ws2.append([n])
        ws2.cell(row=ws2.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws2.column_dimensions["A"].width = 110
    ws3 = wb.create_sheet("Segments")
    _head(ws3, ["Legacy code", "Segment"], [12, 56])
    for new, leg, name, _, _ in K.SEGMENTS:
        ws3.append([leg, name])
    wb.save(path)
    return rows


def write_review_book(path, W):
    sheets = R.review_sheets(W)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Read me"
    lines = [
        ("%s: Early Return Nursing pilot, indemnity verification" % K.KPA, "title"),
        ("Prepared for %s. Valuation date 30 June 2026. Issued %s." % (K.COMPANY, _long(K.REVIEW_DATE)), "sub"),
        ("Each segment sheet lists the segment's sampled claims, both assignment groups together, in claim-number order. "
         "Verification was blind to group.", ""),
        ("Claim number: Haverline's legacy claim number with our control digit after the final hyphen.", ""),
        ("Status: the claim's status at 30 June 2026.", ""),
        ("TTD, TPD, PPD: verified ultimate temporary total, temporary partial and permanent partial disability at "
         "30 June 2026, US dollars. Open claims are projected to ultimate.", ""),
        ("Report: KPA-26-118, %s." % _long(K.REVIEW_DATE), ""),
    ]
    _title_rows(ws, lines)
    ws.column_dimensions["A"].width = 120
    for seg in K.PROG:
        sh = wb.create_sheet(seg)
        _title_rows(sh, [("%s %s" % (seg, K.NAME[seg]), "title"),
                         ("ERN pilot indemnity verification, valuation date 30 June 2026. Verified ultimate, USD.", "sub")])
        hr = _head(sh, ["KPA ref", "Claim number", "Status", "TTD", "TPD", "PPD", "Note"],
                   [11, 16, 8, 12, 11, 12, 44], fill="D7E4E4")
        for row in sheets[seg]:
            sh.append(list(row))
            for col in (4, 5, 6):
                sh.cell(row=sh.max_row, column=col).number_format = "#,##0.00"
        first, last = hr + 1, sh.max_row
        sh.append([])
        sh.append(["Sheet total", None, None, round(sum(r[3] for r in sheets[seg]), 2),
                   round(sum(r[4] for r in sheets[seg]), 2), round(sum(r[5] for r in sheets[seg]), 2), None])
        for col in (4, 5, 6):
            sh.cell(row=sh.max_row, column=col).number_format = "#,##0.00"
            sh.cell(row=sh.max_row, column=col).font = Font(name="Calibri", bold=True)
        sh.cell(row=sh.max_row, column=1).font = Font(name="Calibri", bold=True)
        sh.freeze_panes = sh.cell(row=hr + 1, column=1)
    wb.save(path)
    return sheets


DICT = [
    (K.EXTRACT, "", "One row per claim reported 1 January 2021 to 30 September 2026 whose first report of injury "
                    "recorded time lost from work or that has paid or reserved indemnity. Pulled 30 September 2026."),
    (K.EXTRACT, "claim_no", "Claim number on the claims system (WC and eight digits)."),
    (K.EXTRACT, "legacy_claim_no", "Claim number on the legacy claims system (office letter, year reported, sequence), "
                                   "for claims converted at the 2 June 2025 upgrade."),
    (K.EXTRACT, "office", "Claim office: STL St. Louis, KC Kansas City."),
    (K.EXTRACT, "segment", "Segment code, 2025 coding. Converted claims carry the code their legacy segment converted to."),
    (K.EXTRACT, "class_code", "NCCI classification code of the employer's governing class."),
    (K.EXTRACT, "injury_date", "Date of injury."),
    (K.EXTRACT, "report_date", "Date Haverline received the first report of injury."),
    (K.EXTRACT, "disability_began", "First day of time lost from work on record; blank if none."),
    (K.EXTRACT, "time_off_at_report", "Y when the first report of injury records time lost from work."),
    (K.EXTRACT, "claim_type", "IND when indemnity has been paid or reserved on the claim; MO (medical only) otherwise."),
    (K.EXTRACT, "first_indemnity_date", "Date of the first indemnity payment."),
    (K.EXTRACT, "status", "OPEN, CLOSED or REOPENED at the extract date."),
    (K.EXTRACT, "closed_date", "Date the claim last closed."),
    (K.EXTRACT, "reopened_date", "Date a closed claim was reopened."),
    (K.REGISTER, "", "The triage screen's list of pilot intakes (see the workbook's Notes sheet)."),
    (K.REGISTER, "Claim number", "Legacy claim number."),
    (K.REGISTER, "Segment", "Legacy segment code at triage."),
    (K.REGISTER, "Group", "ERN or Usual handling."),
    (K.NURSE_LOG, "", "ERN case system export: one row per file, and a further row for each reopening of a file. "
                      "Exported 4 August 2025."),
    (K.NURSE_LOG, "case_no", "ERN case number."),
    (K.NURSE_LOG, "claim_ref", "Claim number as entered by the nurse."),
    (K.NURSE_LOG, "team", "ERN pilot team."),
    (K.NURSE_LOG, "nurse", "Nurse case manager's initials."),
    (K.NURSE_LOG, "opened, closed", "Dates the file (or its reopening) opened and closed, MM/DD/YYYY."),
    (K.NURSE_LOG, "closure", "Full duty: worker back at full duty. Back to adjuster: file returned to the adjuster. "
                             "MMI: closed at maximum medical improvement."),
    (K.NURSE_LOG, "reopen", "Y on the row of a reopening."),
    (K.PAYMENTS, "", "Indemnity payments issued 1 January to 30 September 2026, from the payment register. Pulled "
                     "30 September 2026."),
    (K.PAYMENTS, "payment_no", "Payment number."),
    (K.PAYMENTS, "line", "Line of the payment: one line for each week of benefits the payment covers; a settlement has "
                         "one line."),
    (K.PAYMENTS, "benefit", "TTD temporary total, TPD temporary partial, PPD permanent partial disability."),
    (K.PAYMENTS, "payee", "CLAIMANT or ATTORNEY."),
    (K.PAYMENTS, "method", "EFT (ACH direct deposit) or CHK (check)."),
    (K.PAYMENTS, "check_or_trace", "Check number, or the ACH trace number of an EFT."),
    (K.PAYMENTS, "issue_date", "Date the payment was issued (the payment run)."),
    (K.PAYMENTS, "period_from, period_to", "The week of benefits the line covers."),
    (K.PAYMENTS, "amount", "Payment amount, US dollars."),
    (K.PAYMENTS, "status", "ISSUED, or VOID when the check was voided (void_date)."),
    (K.ACH_RETURNS, "", "The bank's ACH return items for Haverline's workers' compensation EFT payments, 1 January to 30 "
                        "September 2026, as received from the bank."),
    (K.CROSSWALK, "", "Legacy segment codes and the 2025 codes they converted to at the 2 June 2025 upgrade."),
]


def write_dictionary(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fields"
    _title_rows(ws, [("Claims data dictionary", "title"),
                     ("%s, Data Engineering. Fields of the claims extract, the pilot register, the ERN case log, the payment "
                      "register and bank return files, and the segment crosswalk in this pack." % K.COMPANY, "sub")])
    _head(ws, ["File", "Field", "Definition"], [40, 22, 100])
    for f, c, d in DICT:
        ws.append([f, c or "(file)", d])
        ws.cell(row=ws.max_row, column=3).alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def write_staffing(path):
    rg = K.rng("staffing")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Positions"
    _title_rows(ws, [("Medical Case Management: 2027 ERN staffing plan", "title"),
                     ("Positions for the seventeen 2027 teams. Status at %s." % _long(dt.date(2026, 10, 30)), "sub")])
    hr = _head(ws, ["Team", "Position", "Role", "Status", "Start", "Grade", "Annual salary midpoint"],
               [7, 10, 32, 15, 11, 7, 20])
    total = 0
    pos = 0
    for t in range(1, K.TEAMS + 1):
        for role, grade, mid in (("Nurse case manager (RN, CCM)", "N4", 92400), ("Return-to-work coordinator", "C2", 58600)):
            pos += 1
            u = rg.random()
            status = "Filled" if u < 0.42 else ("Offer accepted" if u < 0.66 else ("Interviewing" if u < 0.86 else "Posted"))
            start = dt.date(2027, 1, 4) if status in ("Filled", "Offer accepted") else None
            ws.append([t, "MCM-%03d" % (400 + pos), role, status, start, grade, mid])
            ws.cell(row=ws.max_row, column=5).number_format = "yyyy-mm-dd"
            ws.cell(row=ws.max_row, column=7).number_format = "#,##0"
            total += mid
    ws.append([])
    ws.append(["Total", None, None, None, None, None, total])
    ws.cell(row=ws.max_row, column=7).number_format = "#,##0"
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    wb.save(path)


def write_manifest(path, files):
    cols = ["file", "description", "source", "as_of", "prepared_by"]
    K.write_csv(path, cols, [dict(zip(cols, f)) for f in files])
