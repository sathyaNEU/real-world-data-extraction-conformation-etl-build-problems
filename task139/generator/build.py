"""Builds task139's pack into target/ and the golden deliverables into golden/, writes generator/figures.json, and
cuts target.zip; two runs are byte-identical.

    python3 task139/generator/build.py          # the whole build
    python3 task139/generator/build.py pack     # the pack only
"""
import json
import os
import shutil
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
TARGET = os.path.join(TASK, "target")
GOLDEN = os.path.join(TASK, "golden")
sys.path.insert(0, HERE)

import common as K        # noqa: E402


def t(name):
    return os.path.join(TARGET, name)


def pack():
    import openpyxl
    from openpyxl.styles import Font
    import pack as PK
    import world as WD
    for d in (TARGET, GOLDEN):
        shutil.rmtree(d, ignore_errors=True)
        os.makedirs(d)
    C = WD.simulate()
    PK.ids(C)
    n = {}
    n[K.REGISTER] = K.write_csv(t(K.REGISTER), ["contract_number", "vin", "channel", "sale_date", "odometer_at_sale",
                                                "vehicle_class", "contract_price", "lienholder"], PK.register_rows(C))
    n[K.CROSSWALK] = K.write_csv(t(K.CROSSWALK), ["old_contract_number", "new_contract_number"], PK.crosswalk_rows(C))
    n[K.CLAIMS] = K.write_csv(t(K.CLAIMS), ["claim_number", "line_number", "contract_number", "repair_order_date",
                                            "repair_odometer", "facility", "component", "line_type", "amount",
                                            "paid_date"], PK.claim_rows(C), crlf=True)
    n[K.READINGS] = K.write_csv(t(K.READINGS), ["source", "reference", "reading_date", "odometer"], PK.reading_rows(C))
    reqs = PK.cancellation_requests(C)
    n[K.CANCELS] = K.write_csv(t(K.CANCELS), ["request_id", "received_date", "requested_by", "contract_number", "vin",
                                              "lienholder", "odometer_stated", "reason"], PK.cancel_rows(reqs))
    dealers = PK.dealer_rows(C)
    n[K.DEALERS] = K.write_csv(t(K.DEALERS), ["dealer_code", "dealer_name", "city", "state", "dealer_group"], dealers)
    out, returns = PK.remittances(C, dealers)
    n[K.RETURNS] = K.write_csv(t(K.RETURNS), ["trace_number", "originator_name", "amount", "settlement_date",
                                              "return_date", "return_reason"], PK.return_rows(returns, dealers))
    n[K.FACILITIES] = K.write_csv(t(K.FACILITIES), ["facility", "city", "facility_type", "labor_rate_usd_per_hour",
                                                    "in_network"], PK.facility_rows())
    sheets, heads = PK.bordereau_sheets(out)
    wb = openpyxl.Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name
        ws.append(heads[name])
        for c in ws[1]:
            c.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
    wb.save(t(K.BORDEREAUX))
    n[K.BORDEREAUX] = sum(len(v) for v in sheets.values())
    import docs as DOC
    import postprocess as PP
    DOC.write_rate_manual(t(K.RATE_MANUAL))
    DOC.write_terms(t(K.TERMS))
    DOC.write_agreement(t(K.AGREEMENT))
    DOC.write_clip(t(K.CLIP))
    PP.normalize_ooxml(t(K.BORDEREAUX), "Kirsten Weaver", "Kirsten Weaver", "2025-02-03T15:12:00Z",
                       "2026-10-02T13:40:00Z", title="Dealer remittance bordereaux", excel=True,
                       company="Copperline Vehicle Protection")
    for nm, stamp in ((K.RATE_MANUAL, b"D:20260105100000-06'00'"), (K.TERMS, b"D:20230815090000-05'00'"),
                      (K.AGREEMENT, b"D:20260312141500-05'00'"), (K.CLIP, b"D:20250701083000-05'00'")):
        PP.normalize_pdf(t(nm), stamp, stamp)
    return n, C


def cut_zip():
    zp = os.path.join(TASK, "target.zip")
    with zipfile.ZipFile(zp, "w", zipfile.ZIP_DEFLATED) as z:
        for nm in sorted(os.listdir(TARGET)):
            zi = zipfile.ZipInfo(nm, date_time=(2026, 10, 19, 9, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            with open(t(nm), "rb") as f:
                z.writestr(zi, f.read())


def main():
    n, C = pack()
    if len(sys.argv) > 1 and sys.argv[1] == "pack":
        print(json.dumps(n))
        return
    import golden as G
    import goldfiles as GF
    F = G.compute(TARGET)
    json.dump(G.flat(F), open(os.path.join(HERE, "figures.json"), "w"), indent=1, sort_keys=True)
    GF.book(os.path.join(GOLDEN, K.BOOK), F)
    GF.deck(os.path.join(GOLDEN, K.DECK), F)
    import postprocess as PP
    PP.normalize_ooxml(os.path.join(GOLDEN, K.BOOK), "Emily Willis", "Emily Willis", "2026-10-12T15:20:00Z",
                       "2026-10-19T16:45:00Z", title="Drive Plus channel review 2027", excel=True,
                       company="Copperline Vehicle Protection")
    PP.normalize_ooxml(os.path.join(GOLDEN, K.DECK), "Emily Willis", "Nicholas Reed", "2026-10-14T10:05:00Z",
                       "2026-10-19T17:30:00Z", title="Fareline channel renewal 2027", company="Copperline Vehicle Protection")
    import writeup
    writeup.main()
    import metadata
    metadata.main()
    cut_zip()
    print("built: %d files in target/, %d in golden/" % (len(os.listdir(TARGET)), len(os.listdir(GOLDEN))))
    print(json.dumps(n))


if __name__ == "__main__":
    main()
