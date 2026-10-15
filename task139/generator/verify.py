"""The independent verifier: reads only the bundle (target/ and golden/), on its own code path (the csv module and
openpyxl, no pandas, nothing from calc.py or the world), and recomputes the placement, every team saving, the credit,
the default path, and both asks; then checks them against the golden workbook."""
from __future__ import annotations

import csv
import datetime as dt
import math
import os
import re
import sys
from collections import defaultdict

import openpyxl

SEGS = ["HC-SS", "HC-OT", "RW-SS", "RW-OT", "MF-SS", "MF-OT", "CN-SS", "CN-OT"]


def canon(ref):
    t = re.sub(r"[\s-]", "", ref.strip().upper())
    m = re.fullmatch(r"([SK])(\d\d)(\d+)", t)
    return "%s%s-%06d" % (m.group(1), m.group(2), int(m.group(3)))


def rnd(x, step):
    q = x / step
    return int(math.floor(q + 0.5)) * step


def read_csv(path):
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def main(target, golden):
    f = lambda n: os.path.join(target, n)
    names = os.listdir(target)
    pick = lambda pat: [n for n in names if re.search(pat, n)][0]
    cross = {r["legacy_code"]: r["new_code"] for r in read_csv(f(pick(r"crosswalk")))}
    wb = openpyxl.load_workbook(f(pick(r"pilot_assignments")), read_only=True, data_only=True)
    reg = {}
    started = False
    for r in wb["Pilot intakes"].iter_rows(values_only=True):
        if r and r[0] == "Pilot ref":
            started = True
            continue
        if started and r and r[0]:
            reg[canon(r[1])] = dict(seg=cross[r[3]], arm=r[6], office=r[2])
    ext = read_csv(f(pick(r"claims_extract")))
    ctype = {r["legacy_claim_no"]: r["claim_type"] for r in ext if r["legacy_claim_no"]}
    wb = openpyxl.load_workbook(f(pick(r"verification")), read_only=True, data_only=True)
    sev = defaultdict(list)
    for seg in SEGS:
        on = False
        for r in wb[seg].iter_rows(values_only=True):
            if r and r[0] == "KPA ref":
                on = True
                continue
            if not on or not r or not r[0]:
                continue
            if str(r[0]).startswith("Sheet total"):
                break
            leg = canon(str(r[1]).rsplit("-", 1)[0])
            sev[(seg, reg[leg]["arm"])].append(float(r[3]) + float(r[4]) + float(r[5]))
    assert all(len(v) == 100 for v in sev.values()) and len(sev) == 16, {k: len(v) for k, v in sev.items()}
    pdays = defaultdict(int)
    for r in read_csv(f(pick(r"case_log"))):
        o = dt.datetime.strptime(r["opened"], "%m/%d/%Y").date()
        c = dt.datetime.strptime(r["closed"], "%m/%d/%Y").date()
        pdays[canon(r["claim_ref"])] += (c - o).days + 1
    stat = {}
    for seg in SEGS:
        x = {}
        for arm in ("ERN", "Usual handling"):
            ids = [k for k, v in reg.items() if v["seg"] == seg and v["arm"] == arm]
            p = sum(1 for k in ids if ctype[k] == "IND") / len(ids)
            m = sum(sev[(seg, arm)]) / len(sev[(seg, arm)])
            x[arm] = (ids, p, m)
        e_ids = x["ERN"][0]
        L_all = sum(pdays[k] for k in e_ids) / len(e_ids)
        e_ind = [k for k in e_ids if ctype[k] == "IND"]
        L_cl = sum(pdays[k] for k in e_ind) / len(e_ind)
        delta = x["Usual handling"][1] * x["Usual handling"][2] - x["ERN"][1] * x["ERN"][2]
        d = x["Usual handling"][2] - x["ERN"][2]
        stat[seg] = dict(V=7300 * delta / L_all, Vd=7300 * d / L_cl, L_all=L_all, L_cl=L_cl)
    vol = defaultdict(int)
    for r in ext:
        if "2025-10-01" <= r["report_date"] <= "2026-09-30" and r["time_off_at_report"] == "Y" and r["segment"] in SEGS:
            assert not r["legacy_claim_no"]
            vol[r["segment"]] += 1

    def place(key, Lkey):
        teams, left = {}, 17
        for s in sorted(SEGS, key=lambda s: (-stat[s][key], s)):
            cap = min(3, int(math.floor(vol[s] * stat[s][Lkey] / 7300 + 1e-12)))
            teams[s] = min(3, cap, left)
            left -= teams[s]
        return teams, sum(teams[s] * stat[s][key] for s in SEGS)
    teams, credit = place("V", "L_all")
    dteams, dcredit = place("Vd", "L_cl")
    assert rnd(credit, 100000) == 4100000, credit
    assert rnd(dcredit, 100000) == 2000000 and dteams["HC-SS"] == 0 and dteams["RW-SS"] == 0, (dteams, dcredit)
    # ask 1: intakes reported in 2025, the converted rows read from their dates, the parallel-run copies once
    key = lambda r: (r["office"], r["segment"], r["class_code"], r["injury_date"], r["report_date"], r["disability_began"])
    natives = set(key(r) for r in ext if not r["legacy_claim_no"])
    a1 = defaultdict(int)
    for r in ext:
        if r["segment"] not in SEGS or not ("2025-01-01" <= r["report_date"] <= "2025-12-31"):
            continue
        if r["legacy_claim_no"]:
            took = r["disability_began"] != "" and r["disability_began"] <= r["report_date"]
            if "2025-05-29" <= r["report_date"] <= "2025-06-01" and key(r) in natives:
                continue
        else:
            took = r["time_off_at_report"] == "Y"
        a1[r["segment"]] += took
    # ask 2: each payment once, issued, not returned by the bank
    seg_of = {r["claim_no"]: r["segment"] for r in ext}
    returned = {r["trace_number"] for r in read_csv(f(pick(r"ach_returns")))}
    seen, a2 = set(), defaultdict(float)
    for r in read_csv(f(pick(r"indemnity_payments"))):
        if r["payment_no"] in seen:
            continue
        seen.add(r["payment_no"])
        if r["status"] != "ISSUED" or r["check_or_trace"] in returned or seg_of[r["claim_no"]] not in SEGS:
            continue
        a2[seg_of[r["claim_no"]]] += float(r["amount"])
    gb = openpyxl.load_workbook(os.path.join(golden, [n for n in os.listdir(golden) if n.endswith(".xlsx")][0]),
                                data_only=True)
    pl = {r[0]: (r[3], r[4]) for r in gb["Placement"].iter_rows(values_only=True) if r and r[0] in SEGS}
    for s in SEGS:
        assert pl[s] == (teams[s], rnd(stat[s]["V"], 10000)), (s, pl[s], teams[s], stat[s]["V"])
    tot = [r for r in gb["Placement"].iter_rows(values_only=True) if r and r[0] and str(r[0]).startswith("AY2027")][0][4]
    assert tot == rnd(credit, 100000), (tot, credit)
    g1 = {r[0]: r[1] for r in gb["2025 intakes"].iter_rows(values_only=True) if r and r[0] in SEGS}
    g2 = {r[0]: r[1] for r in gb["Indemnity paid 2026"].iter_rows(values_only=True) if r and r[0] in SEGS}
    assert all(g1[s] == a1[s] for s in SEGS), (g1, dict(a1))
    assert all(g2[s] == rnd(a2[s], 1000) for s in SEGS), (g2, dict(a2))
    print("verify.py: placement %s, credit %d (unrounded %.0f); default path %d; asks reproduce" %
          (" ".join("%s:%d" % (s, teams[s]) for s in SEGS), rnd(credit, 100000), credit, rnd(dcredit, 100000)))
    return 0


if __name__ == "__main__":
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    sys.exit(main(os.path.join(here, "target"), os.path.join(here, "golden")))
