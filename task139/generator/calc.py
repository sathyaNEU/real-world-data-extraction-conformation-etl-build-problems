"""The golden and every rung, computed from the shipped pack alone (the files under target/), the way an analyst
would read them: the register for the pilot's intakes and groups, the claims extract for which of them became
indemnity claims and for the units' intakes, Kellerman Pryce's workbook for the verified amounts, the case log for the
place-days a file held, the charter and the pick memo for the rules. Nothing here reads the generator's world."""
from __future__ import annotations

import datetime as dt
import math
import os
import re

import numpy as np
import openpyxl
import pandas as pd

import common as K


def norm_legacy(ref):
    """A legacy claim number in its canonical form, S23-000740, from the ways the systems and the nurses write it."""
    t = re.sub(r"[\s-]", "", str(ref).strip().upper())
    m = re.fullmatch(r"([SK])(\d{2})(\d+)", t)
    if not m:
        raise ValueError(ref)
    return "%s%s-%06d" % (m.group(1), m.group(2), int(m.group(3)))


def strip_check(no):
    head, tail = str(no).rsplit("-", 1)
    assert tail.isdigit() and len(tail) == 1, no
    return head


def load(target):
    P = lambda f: os.path.join(target, f)
    D = {}
    D["extract"] = pd.read_csv(P(K.EXTRACT), dtype=str, keep_default_na=False)
    D["cross"] = pd.read_csv(P(K.CROSSWALK), dtype=str, keep_default_na=False)
    wb = openpyxl.load_workbook(P(K.REGISTER), read_only=True, data_only=True)
    ws = wb["Pilot intakes"]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "Pilot ref")
    reg = pd.DataFrame([r for r in rows[hi + 1:] if r and r[0]], columns=list(rows[hi]))
    leg2new = dict(zip(D["cross"]["legacy_code"], D["cross"]["new_code"]))
    reg["seg"] = reg["Segment"].map(leg2new)
    reg["legacy"] = reg["Claim number"].map(norm_legacy)
    reg["arm"] = reg["Group"].map({"ERN": "E", "Usual handling": "U"})
    D["register"] = reg
    wb = openpyxl.load_workbook(P(K.REVIEW_BOOK), read_only=True, data_only=True)
    rv = []
    for seg in K.PROG:
        ws = wb[seg]
        rows = list(ws.iter_rows(values_only=True))
        hi = next(i for i, r in enumerate(rows) if r and r[0] == "KPA ref")
        for r in rows[hi + 1:]:
            if not r or not r[0] or str(r[0]).startswith("Sheet total"):
                continue
            rv.append(dict(seg=seg, ref=r[0], legacy=norm_legacy(strip_check(r[1])), status=r[2], ttd=float(r[3]),
                           tpd=float(r[4]), ppd=float(r[5])))
    D["review"] = pd.DataFrame(rv)
    log = pd.read_csv(P(K.NURSE_LOG), dtype=str, keep_default_na=False)
    log["legacy"] = log["claim_ref"].map(norm_legacy)
    log["o"] = pd.to_datetime(log["opened"], format="%m/%d/%Y")
    log["c"] = pd.to_datetime(log["closed"], format="%m/%d/%Y")
    log["pd"] = (log["c"] - log["o"]).dt.days + 1          # charter 5.3: the opening and closing days both count
    D["log"] = log
    D["pay"] = pd.read_csv(P(K.PAYMENTS), dtype=str, keep_default_na=False)
    D["ach"] = pd.read_csv(P(K.ACH_RETURNS), dtype=str, keep_default_na=False)
    return D


def pilot_metrics(D, reopen=True, inclusive=True, ind_from="claim_type", office=None):
    """Per segment: each arm's intakes and indemnity claims, the review's verified mean indemnity per sampled claim, the
    place-days an ERN intake's file held and an ERN indemnity claim's file held."""
    reg, ex, rv, log = D["register"], D["extract"], D["review"], D["log"]
    conv = ex[ex["legacy_claim_no"] != ""]
    if ind_from == "claim_type":
        ind = dict(zip(conv["legacy_claim_no"], conv["claim_type"] == "IND"))
    else:
        ind = dict(zip(conv["legacy_claim_no"], conv["first_indemnity_date"] != ""))
    reg = reg.assign(ind=reg["legacy"].map(ind))
    assert reg["ind"].notna().all()
    lg = log if reopen else log[log["reopen"] != "Y"]
    days = lg["pd"] if inclusive else lg["pd"] - 1
    pdays = days.groupby(lg["legacy"]).sum()
    reg = reg.assign(pdays=reg["legacy"].map(pdays))
    rvj = rv.merge(reg[["legacy", "arm"]], on="legacy", how="left")
    assert rvj["arm"].notna().all()
    rvj["tot"] = rvj["ttd"] + rvj["tpd"] + rvj["ppd"]
    out = {}
    for seg in K.PROG:
        r = reg[reg["seg"] == seg]
        if office:
            r = r[r["Office"] == office]
        m = {}
        for arm in ("E", "U"):
            ra = r[r["arm"] == arm]
            s = rvj[(rvj["seg"] == seg) & (rvj["arm"] == arm)]
            m[arm] = dict(n=len(ra), n_ind=int(ra["ind"].sum()), p=float(ra["ind"].mean()), sev=float(s["tot"].mean()),
                          n_s=len(s))
        e = r[r["arm"] == "E"]
        assert e["pdays"].notna().all()
        L_all = float(e["pdays"].mean())
        L_cl = float(e[e["ind"]]["pdays"].mean())
        es = rvj[(rvj["seg"] == seg) & (rvj["arm"] == "E")]
        L_sam = float(e[e["legacy"].isin(es["legacy"])]["pdays"].mean())
        delta = m["U"]["p"] * m["U"]["sev"] - m["E"]["p"] * m["E"]["sev"]
        d = m["U"]["sev"] - m["E"]["sev"]
        out[seg] = dict(nU=m["U"]["n"], nE=m["E"]["n"], indU=m["U"]["n_ind"], indE=m["E"]["n_ind"], pU=m["U"]["p"],
                        pE=m["E"]["p"], sevU=m["U"]["sev"], sevE=m["E"]["sev"], L_all=L_all, L_cl=L_cl, L_sam=L_sam,
                        delta=delta, d=d)
    return out


def intakes(D, a, b, fix_vintage=True, dedup=True, by="report_date", strict=False):
    """Each segment's intakes reported a to b: a first report recording time lost (the dictionary: disability_began on
    or before report_date). Rows converted at the 2 June 2025 upgrade carry the legacy LOST_TIME indicator in
    time_off_at_report, so for them the dates decide; a claim keyed in both systems in the parallel run is listed twice
    and counted once. The rivals: fix_vintage=False reads the flag as filed; strict counts time lost beginning before
    the report date only; dedup=False keeps both listings, "window" drops every converted claim of the parallel run,
    "global" drops every repeat of the six fields across the period."""
    ex = D["extract"]
    x = ex[(ex[by] >= a.isoformat()) & (ex[by] <= b.isoformat()) & ex["segment"].isin(K.PROG)]
    conv = x["legacy_claim_no"] != ""
    key = ["office", "segment", "class_code", "injury_date", "report_date", "disability_began"]
    a0, b0 = K.PARALLEL
    if fix_vintage:
        took = (x["disability_began"] != "") & ((x["disability_began"] < x["report_date"]) if strict else
                                                (x["disability_began"] <= x["report_date"]))
        flag = np.where(conv, took, x["time_off_at_report"] == "Y")
    else:
        flag = (x["time_off_at_report"] == "Y").values
    x = x[flag]
    if dedup == "window":
        x = x[~((x["legacy_claim_no"] != "") & (x["report_date"] >= a0.isoformat()) & (x["report_date"] <= b0.isoformat()))]
    elif dedup == "global":
        x = x.drop_duplicates(subset=key)
    elif dedup:
        nat = x[x["legacy_claim_no"] == ""]
        natkeys = set(map(tuple, nat[key].values.tolist()))
        cv = x[(x["legacy_claim_no"] != "") & (x["report_date"] >= a0.isoformat()) & (x["report_date"] <= b0.isoformat())]
        dup = cv[[tuple(r) in natkeys for r in cv[key].values.tolist()]]
        x = x.drop(dup.index)
    return x.groupby("segment").size().reindex(K.PROG, fill_value=0).to_dict()


def caps(L, vol):
    return {s: min(K.MAX_TEAMS, int(math.floor(vol[s] * L[s] / K.TEAM_PLACE_DAYS + 1e-12))) for s in K.PROG}


def allocate(value, cap=None):
    order = sorted(K.PROG, key=lambda s: (-value[s], s))
    teams, left = {s: 0 for s in K.PROG}, K.TEAMS
    for s in order:
        t = min(K.MAX_TEAMS, (cap or {}).get(s, K.MAX_TEAMS), left)
        teams[s] = t
        left -= t
    return teams


def rung(M, saving="delta", L="L_all", cap_L="L_all", vol=None, rank="value"):
    """One complete pipeline: each segment's saving per team-year (7,300 place-days over the place-days that unit of
    saving holds), the fill cap from the unit's intakes, the allocation and the credit."""
    value = {s: K.TEAM_PLACE_DAYS * M[s][saving] / M[s][L] for s in K.PROG}
    cap = caps({s: M[s][cap_L] for s in K.PROG}, vol) if (vol is not None and cap_L) else None
    key = value if rank == "value" else {s: M[s][saving] for s in K.PROG}
    teams = allocate(key, cap)
    credit = sum(teams[s] * value[s] for s in K.PROG)
    return dict(value=value, teams=teams, credit=credit, cap=cap)


def paid_2026(D, dedup_lines=True, drop_returns=True, drop_void=True):
    pay, ach, ex = D["pay"], D["ach"], D["extract"]
    seg = dict(zip(ex["claim_no"], ex["segment"]))
    x = pay.copy()
    if dedup_lines:
        x = x[x["line"] == "1"]
    if drop_void:
        x = x[x["status"] == "ISSUED"]
    if drop_returns:
        x = x[~x["check_or_trace"].isin(set(ach["trace_number"]))]
    x = x.assign(seg=x["claim_no"].map(seg), amt=x["amount"].astype(float))
    x = x[x["seg"].isin(K.PROG)]
    return x.groupby("seg")["amt"].sum().reindex(K.PROG, fill_value=0.0).to_dict()


def golden(target):
    D = load(target)
    M = pilot_metrics(D)
    for s in K.PROG:
        M[s]["V"] = K.TEAM_PLACE_DAYS * M[s]["delta"] / M[s]["L_all"]
    v12 = intakes(D, dt.date(2025, 10, 1), K.EXTRACT_END)
    g = rung(M, vol=v12)
    a1 = intakes(D, dt.date(2025, 1, 1), dt.date(2025, 12, 31))
    a2 = paid_2026(D)
    return dict(D=D, M=M, v12=v12, g=g, ask1=a1, ask2=a2)
