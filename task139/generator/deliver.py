"""The golden deliverables, computed from the shipped pack by calc.py: the placement workbook and the one-page note
with its chart; and generator/figures.json, every figure the submission states."""
from __future__ import annotations

import datetime as dt
import io
import json
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import calc as C
import common as K

ORDER_NAME = {"HC-SS": "healthcare and social assistance strains and sprains",
              "HC-OT": "healthcare and social assistance other injuries",
              "RW-SS": "retail and warehousing strains and sprains", "RW-OT": "retail and warehousing other injuries",
              "MF-SS": "manufacturing strains and sprains", "MF-OT": "manufacturing other injuries",
              "CN-SS": "construction strains and sprains", "CN-OT": "construction other injuries"}


def figures(G):
    M, g = G["M"], G["g"]
    f = {"teams": {s: g["teams"][s] for s in K.PROG},
         "team_saving": {s: K.round_to(g["value"][s], 10000) for s in K.PROG},
         "team_saving_exact": {s: round(g["value"][s], 2) for s in K.PROG},
         "credit": K.round_to(g["credit"], 100000), "credit_exact": round(g["credit"], 2),
         "ask1_intakes_2025": {s: int(G["ask1"][s]) for s in K.PROG},
         "ask2_paid_2026": {s: K.round_to(G["ask2"][s], 1000) for s in K.PROG},
         "ask2_paid_2026_exact": {s: round(G["ask2"][s], 2) for s in K.PROG},
         "v12": {s: int(G["v12"][s]) for s in K.PROG}, "cap": g["cap"],
         "pilot": {s: {k: (round(v, 6) if isinstance(v, float) else v) for k, v in M[s].items()} for s in K.PROG}}
    f["pilot_intakes"] = int(sum(M[s]["nU"] + M[s]["nE"] for s in K.PROG))
    f["pilot_paid_indemnity"] = int(sum(M[s]["indU"] + M[s]["indE"] for s in K.PROG))
    f["sampled_per_group"] = 100
    f["places"] = K.TEAMS * K.PLACES
    f["team_place_days"] = K.TEAM_PLACE_DAYS
    # the critical components as the write-up states them
    f["components"] = {
        "saving_per_intake": {s: "%.0f" % M[s]["delta"] for s in K.PROG},
        "paid_share_ern_pct": {s: "%.1f" % (100 * M[s]["pE"]) for s in K.PROG},
        "paid_share_uh_pct": {s: "%.1f" % (100 * M[s]["pU"]) for s in K.PROG},
        "place_days_per_ern_intake": {s: "%.2f" % M[s]["L_all"] for s in K.PROG},
        "fill_teams_worth": {s: "%.2f" % (G["v12"][s] * M[s]["L_all"] / K.TEAM_PLACE_DAYS) for s in K.PROG},
    }
    return f


def _fmt_money(x):
    return "$%s" % format(int(x), ",")


def write_book(path, G, F):
    M, g = G["M"], G["g"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Placement"
    ws.append(["Early Return Nursing: 2027 placement and AY2027 indemnity credit"])
    ws["A1"].font = Font(name="Calibri", bold=True, size=13)
    ws.append(["Seventeen teams of twenty places; at most three to a claim unit and no more than a unit's intakes keep "
               "in use. Saving per team-year at run rate (7,300 place-days), verified ultimate indemnity at 30 June 2026."])
    ws["A2"].font = Font(name="Calibri", italic=True, size=10, color="595959")
    ws.append([])
    head = ["Code", "Segment", "Claim unit", "Teams 2027", "Saving per team-year (USD)"]
    ws.append(head)
    hr = ws.max_row
    for c in range(1, len(head) + 1):
        cell = ws.cell(row=hr, column=c)
        cell.font = Font(name="Calibri", bold=True)
        cell.fill = PatternFill("solid", fgColor="DCE3EC")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    order = sorted(K.PROG, key=lambda s: -g["value"][s])
    for s in order:
        ws.append([s, K.NAME[s], K.UNIT[s], F["teams"][s], F["team_saving"][s]])
        ws.cell(row=ws.max_row, column=5).number_format = "#,##0"
    ws.append([])
    ws.append(["Total", None, None, sum(F["teams"].values()), None])
    ws.append(["AY2027 indemnity credit at run rate (USD)", None, None, None, F["credit"]])
    ws.cell(row=ws.max_row, column=5).number_format = "#,##0"
    for r in (ws.max_row - 1, ws.max_row):
        ws.cell(row=r, column=1).font = Font(name="Calibri", bold=True)
    for i, w in enumerate([10, 52, 26, 11, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)

    b = wb.create_sheet("Pilot basis")
    b.append(["ERN pilot read by assigned intake (triage dates 1 March 2023 to 28 February 2025)"])
    b["A1"].font = Font(name="Calibri", bold=True, size=12)
    b.append(["Indemnity per intake is the group's share of intakes that paid indemnity times its mean verified "
              "ultimate indemnity per sampled indemnity claim. Place-days per ERN intake from the case log, every file "
              "and reopening, opening and closing days counted."])
    b["A2"].font = Font(name="Calibri", italic=True, size=10, color="595959")
    b.append([])
    cols = ["Segment", "Intakes, usual handling", "Intakes, ERN", "Intakes paying indemnity, usual handling",
            "Intakes paying indemnity, ERN", "Mean verified indemnity per claim, usual handling (USD)",
            "Mean verified indemnity per claim, ERN (USD)", "Indemnity per intake, usual handling (USD)",
            "Indemnity per intake, ERN (USD)", "Saving per intake (USD)", "Place-days per ERN intake",
            "Intakes Oct 2025 to Sep 2026", "Teams' places those intakes fill"]
    b.append(cols)
    hb = b.max_row
    for c in range(1, len(cols) + 1):
        cell = b.cell(row=hb, column=c)
        cell.font = Font(name="Calibri", bold=True)
        cell.fill = PatternFill("solid", fgColor="DCE3EC")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for s in K.PROG:
        m = M[s]
        b.append([s, m["nU"], m["nE"], m["indU"], m["indE"], round(m["sevU"], 2), round(m["sevE"], 2),
                  round(m["pU"] * m["sevU"], 2), round(m["pE"] * m["sevE"], 2), round(m["delta"], 2),
                  round(m["L_all"], 2), F["v12"][s], float(F["components"]["fill_teams_worth"][s])])
        for c in range(6, 11):
            b.cell(row=b.max_row, column=c).number_format = "#,##0.00"
        b.cell(row=b.max_row, column=11).number_format = "0.00"
        b.cell(row=b.max_row, column=12).number_format = "#,##0"
        b.cell(row=b.max_row, column=13).number_format = "0.00"
    for i, w in enumerate([10] + [13] * 12, 1):
        b.column_dimensions[get_column_letter(i)].width = w
    b.row_dimensions[hb].height = 75
    b.freeze_panes = b.cell(row=hb + 1, column=2)

    a1 = wb.create_sheet("2025 intakes")
    a1.append(["Intakes reported in 2025 by segment"])
    a1["A1"].font = Font(name="Calibri", bold=True, size=12)
    a1.append(["A claim whose first report of injury records time lost from work, counted once."])
    a1["A2"].font = Font(name="Calibri", italic=True, size=10, color="595959")
    a1.append([])
    a1.append(["Segment", "Intakes 2025"])
    for c in (1, 2):
        a1.cell(row=a1.max_row, column=c).font = Font(name="Calibri", bold=True)
    for s in K.PROG:
        a1.append([s, F["ask1_intakes_2025"][s]])
        a1.cell(row=a1.max_row, column=2).number_format = "#,##0"
    a1.column_dimensions["A"].width = 12
    a1.column_dimensions["B"].width = 14

    a2 = wb.create_sheet("Indemnity paid 2026")
    a2.append(["Indemnity paid on each segment's claims, January to September 2026"])
    a2["A1"].font = Font(name="Calibri", bold=True, size=12)
    a2.append(["Temporary total, temporary partial and permanent partial disability; each payment once; voided checks "
               "and EFTs the bank returned left out (a returned EFT's re-issued check counts)."])
    a2["A2"].font = Font(name="Calibri", italic=True, size=10, color="595959")
    a2.append([])
    a2.append(["Segment", "Indemnity paid (USD)"])
    for c in (1, 2):
        a2.cell(row=a2.max_row, column=c).font = Font(name="Calibri", bold=True)
    for s in K.PROG:
        a2.append([s, F["ask2_paid_2026"][s]])
        a2.cell(row=a2.max_row, column=2).number_format = "#,##0"
    a2.column_dimensions["A"].width = 12
    a2.column_dimensions["B"].width = 20

    n = wb.create_sheet("Notes")
    notes = [
        "Saving per intake: the pilot assigned intakes at random at triage, so each group is read per assigned intake. "
        "Nurses kept part of the strain intakes at work on suitable duties and those intakes never paid indemnity; the "
        "comparison of verified amounts per indemnity claim alone sets ERN's remaining claims against all of usual "
        "handling's and misses that saving.",
        "Saving per team-year: a team turns its twenty places over at the place-days an ERN intake held in the pilot "
        "(7,300 place-days a year at run rate), so a team's year is 7,300 times the saving per intake over the "
        "place-days per intake.",
        "Teams: highest saving per team-year first, three to a unit, and no more than the unit's intakes keep in use "
        "(intakes in the twelve months to 30 September 2026 times place-days per intake over 7,300). Retail and "
        "warehousing strains and sprains keeps two teams' places in use, not three.",
        "Amounts: Kellerman Pryce's verified ultimate indemnity at 30 June 2026, no trend. Indemnity only.",
        "2025 intakes: claims converted at the June 2025 upgrade carry the legacy lost-time indicator in "
        "time_off_at_report; for them the first report's own dates decide. Claims keyed in both systems in the "
        "29 May to 1 June 2025 parallel run are counted once.",
        "Indemnity paid: the payment register repeats a payment's amount on each weekly line, so each payment is counted "
        "once; voided checks and the EFTs in the bank's return file are left out.",
        "Rounding: team savings to the nearest $10,000 and the credit to the nearest $100,000, from the unrounded team "
        "savings; indemnity paid to the nearest $1,000.",
        "As of: claims extract and payment register pulled 30 September 2026; Kellerman Pryce valuation 30 June 2026. "
        "Prepared by Medical Case Management, 20 November 2026.",
    ]
    n.append(["Notes"])
    n["A1"].font = Font(name="Calibri", bold=True, size=12)
    for t in notes:
        n.append([t])
        n.cell(row=n.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    n.column_dimensions["A"].width = 120
    wb.save(path)


def chart_png(G, F):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import rcParams
    rcParams["font.family"] = "DejaVu Sans"
    g = G["g"]
    order = sorted(K.PROG, key=lambda s: -g["value"][s])
    vals = [F["team_saving"][s] / 1000.0 for s in order]
    teams = [F["teams"][s] for s in order]
    fill = {3: "#1f4e79", 2: "#6f9fcb", 0: "#c3c9d1"}
    fig, ax = plt.subplots(figsize=(7.0, 3.1), dpi=160)
    bars = ax.bar(range(len(order)), vals, color=[fill[t] for t in teams], width=0.62, edgecolor="white", linewidth=1.0)
    for b, t in zip(bars, teams):
        ax.text(b.get_x() + b.get_width() / 2, b.get_height() + 5, ("%d teams" % t) if t else "no team",
                ha="center", va="bottom", fontsize=7.6, color="#1d2330")
    ax.set_xticks(range(len(order)))
    ax.set_xticklabels(["%s\n%s" % (s, K.SHORT[s].split(" ")[0].lower().replace("manufacturing", "manuf.")
                                     .replace("construction", "constr.").replace("healthcare", "health").replace(
                                         "retail", "retail") + (" strains" if s.endswith("SS") else " other"))
                        for s in order], fontsize=7.0, color="#3a4150")
    ax.set_ylabel("Saving per team-year ($000)", fontsize=8, color="#3a4150")
    ax.set_ylim(0, 340)
    ax.set_yticks([0, 100, 200, 300])
    ax.tick_params(axis="y", labelsize=7.5, colors="#3a4150", length=0)
    ax.tick_params(axis="x", length=0)
    ax.grid(axis="y", color="#e3e6ea", linewidth=0.6)
    ax.set_axisbelow(True)
    for sp in ("top", "right", "left"):
        ax.spines[sp].set_visible(False)
    ax.spines["bottom"].set_color("#9aa2ad")
    ax.set_title("A strain team saves the most; retail strains' intakes fill two teams, not three",
                 fontsize=9, loc="left", color="#1d2330")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def write_note(path, G, F):
    from reportlab import rl_config
    rl_config.invariant = 1
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    import docs as DOC
    g = G["g"]
    body = ParagraphStyle("b", fontName=DOC.SERIF, fontSize=9.8, leading=12.8, textColor=DOC.INK, spaceAfter=5)
    head = ParagraphStyle("h", fontName=DOC.SANS_B, fontSize=12.6, leading=15.4, textColor=DOC.NAVY, spaceAfter=4)
    memo = ParagraphStyle("m", fontName=DOC.SANS, fontSize=9.2, leading=11.6, textColor=DOC.INK)
    foot = ParagraphStyle("f", fontName=DOC.SANS_I, fontSize=7.8, leading=9.8, textColor=DOC.GREY, spaceBefore=3)
    three = [s for s in sorted(K.PROG, key=lambda s: -g["value"][s]) if F["teams"][s] == 3]
    two = [s for s in K.PROG if F["teams"][s] == 2]
    none = [s for s in K.PROG if F["teams"][s] == 0]
    lst = lambda ss: ", ".join("%s (%s)" % (ORDER_NAME[s], s) for s in ss)
    credit_m = "%.1f" % (F["credit"] / 1e6)
    s = [Paragraph("<b>To</b> Executive Committee&nbsp;&nbsp;&nbsp;<b>From</b> %s, Director, Medical Case Management"
                   % K.PEOPLE["director"], memo),
         Paragraph("<b>Copy</b> %s, Chief Pricing Actuary&nbsp;&nbsp;&nbsp;<b>Date</b> 20 November 2026"
                   % K.PEOPLE["pricing"], memo),
         Paragraph("<b>Re</b> 2027 plan: Early Return Nursing placement and the AY2027 indemnity credit", memo),
         Spacer(1, 8),
         Paragraph("Seventeen ERN teams: three each to five units, two to retail strains, a $%s million AY2027 credit"
                   % credit_m, head),
         Paragraph("I propose we place three Early Return Nursing teams each in %s, and two in %s; none in %s. At run "
                   "rate the seventeen teams save $%s of AY2027 ultimate indemnity, the credit Pricing books in the "
                   "pick.<super>1</super>" % (lst(three), lst(two), " or ".join(ORDER_NAME[x] for x in none),
                                              format(F["credit"], ",")), body)]
    png = chart_png(G, F)
    s.append(Image(io.BytesIO(png), width=6.6 * inch, height=6.6 * inch * 3.1 / 7.0))
    s.append(Spacer(1, 4))
    rows = [["Unit", "Segment", "Teams", "Saving per team-year"]]
    for x in sorted(K.PROG, key=lambda x: -g["value"][x]):
        rows.append([K.UNIT[x], "%s %s" % (x, ORDER_NAME[x]), str(F["teams"][x]), "$%s" % format(F["team_saving"][x], ",")])
    rows.append(["Total", "", str(sum(F["teams"].values())), "$%s credit" % format(F["credit"], ",")])
    t = Table(rows, colWidths=[1.75 * inch, 3.05 * inch, 0.55 * inch, 1.35 * inch], hAlign="LEFT")
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), DOC.SANS, 8.4), ("FONT", (0, 0), (-1, 0), DOC.SANS_B, 8.4),
                           ("FONT", (0, -1), (-1, -1), DOC.SANS_B, 8.4), ("TEXTCOLOR", (0, 0), (-1, -1), DOC.INK),
                           ("LINEBELOW", (0, 0), (-1, 0), 0.6, DOC.RULE), ("LINEABOVE", (0, -1), (-1, -1), 0.6, DOC.RULE),
                           ("ALIGN", (2, 0), (-1, -1), "RIGHT"), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6),
                           ("TOPPADDING", (0, 0), (-1, -1), 1.2)]))
    s.append(t)
    s.append(Spacer(1, 6))
    s.append(Paragraph(
        "The pilot assigned intakes to the nurses at random at triage, and in the strain segments the nurses kept part of "
        "their intakes at work on suitable duties, so those intakes never paid indemnity: in healthcare strains %s%% of "
        "the nurses' intakes paid indemnity against %s%% under usual handling. Read by assigned intake, as the pilot "
        "was randomised, a healthcare strain intake saves $%s, and its file is short (%s place-days), so a team turns "
        "its twenty places over often and saves $%s a year. Read on the verified indemnity claims alone, the strain "
        "segments look like a few per cent of a claim and the healthcare and retail strain units would get no team."
        % (F["components"]["paid_share_ern_pct"]["HC-SS"], F["components"]["paid_share_uh_pct"]["HC-SS"],
                   F["components"]["saving_per_intake"]["HC-SS"], F["components"]["place_days_per_ern_intake"]["HC-SS"],
                   format(F["team_saving"]["HC-SS"], ",")), body))
    s.append(Paragraph(
        "Retail and warehousing strains would take a third team on saving alone, but its %s intakes in the twelve months "
        "to September fill %s teams' places, so the seventeenth team goes to healthcare other injuries."
        % (format(F["v12"]["RW-SS"], ","), F["components"]["fill_teams_worth"]["RW-SS"]), body))
    s.append(Paragraph("<super>1</super> At run rate, every place in use every day of 2027, at Kellerman Pryce's "
                       "verified ultimate indemnity at 30 June 2026 with no trend (AY2027 pick memo, 28 October 2026). "
                       "Indemnity only; medical sits in the medical pick.", foot))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(DOC.SANS, 7.4)
        canvas.setFillColor(DOC.GREY)
        canvas.drawString(0.9 * inch, 0.5 * inch, "%s  |  Medical Case Management" % K.COMPANY)
        canvas.drawRightString(LETTER[0] - 0.9 * inch, 0.5 * inch, "MCM-2027-04  |  Page %d" % doc.page)
        canvas.restoreState()
    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=0.9 * inch, rightMargin=0.9 * inch, topMargin=0.6 * inch,
                            bottomMargin=0.7 * inch, title="ERN placement 2027", author=K.PEOPLE["director"],
                            subject="", creator="", producer="")
    doc.build(s, onFirstPage=footer, onLaterPages=footer)


def main(target, golden, figures_path):
    G = C.golden(target)
    F = figures(G)
    os.makedirs(golden, exist_ok=True)
    write_book(os.path.join(golden, K.GOLD_BOOK), G, F)
    write_note(os.path.join(golden, K.GOLD_NOTE), G, F)
    import postprocess as PP
    PP.normalize_ooxml(os.path.join(golden, K.GOLD_BOOK), K.PEOPLE["director"], K.PEOPLE["director"],
                       "2026-11-16T15:02:44Z", "2026-11-20T21:37:10Z", excel=True)
    PP.normalize_pdf(os.path.join(golden, K.GOLD_NOTE), b"D:20261120154012-06'00'", b"D:20261120154512-06'00'")
    with open(figures_path, "w") as fh:
        json.dump(F, fh, indent=1, sort_keys=True)
    return G, F


if __name__ == "__main__":
    sys.exit(main(sys.argv[1], sys.argv[2], sys.argv[3]) and 0)
