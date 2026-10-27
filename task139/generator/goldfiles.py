"""The golden deliverables: the product committee's renewal deck and the channel review workbook."""
from __future__ import annotations

import io
import re
import zipfile

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import asks as AK
import common as K

INK = "1D2330"
RUST = "7A4B1E"


def r5(x):
    return int(5 * round(x / 5.0))


def _sheet(ws, title, header, rows, widths, fmts=None):
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=12, color=RUST)
    ws.append([])
    ws.append(header)
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=RUST)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(r)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A4"
    for col, f in (fmts or {}).items():
        for row in ws.iter_rows(min_row=4, min_col=col, max_col=col):
            for c in row:
                if isinstance(c.value, (int, float)):
                    c.number_format = f
                    c.alignment = Alignment(horizontal="right")


def book(path, F):
    g = F['g']
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Renewal forecast"
    _sheet(ws, "Fareline channel: lifetime loss ratio of 2027 sales (rate manual section 5)",
           ["Item", "Value"],
           [["Election for 1 January 2027", g['call']],
            ["Lifetime loss ratio at the current price (nearest 5 per cent)", "%d%%" % r5(g['lr'])],
            ["Lifetime loss ratio with the 40 per cent high-use surcharge (nearest 5 per cent)", "%d%%" % r5(g['lr_s'])],
            ["Claims cost over a 2027 contract's cover (average, USD)", round(g['cost'], 0)],
            ["Contract price at current rates (average, USD)", round(g['price'], 0)],
            ["2026 Fareline contracts the forecast is built on", g['n']]],
           [70, 40], {2: "#,##0"})
    B = wb.create_sheet("Cost per mile by odometer")
    rows = []
    for j in range(len(g['cpm'])):
        if g['miles'][j] <= 0:
            continue
        lo = int(g['edges'][j])
        rows.append(["%s to %s" % (format(lo, ","), format(lo + 9999, ",")), round(g['miles'][j] / 1e6, 2),
                     round(g['cpm'][j] * 1000, 2)])
    _sheet(B, "Claims cost per 1,000 miles driven under cover, by odometer band, whole book 2019 to September 2026",
           ["Odometer band (miles)", "Miles under cover (millions)", "Claims cost per 1,000 miles (USD)"], rows,
           [24, 22, 26], {2: "0.00", 3: "0.00"})
    A = wb.create_sheet("Cancellations 2026")
    _sheet(A, "Retail contracts cancelled in 2026 and the refunds owed (contract terms section 5)",
           ["Quarter", "Contracts cancelled", "Refunds owed (USD)", "Of which to lienholders (USD)"],
           [[q, F['ask_a'][q]['n'], round(F['ask_a'][q]['refund'], 0), round(F['ask_a'][q]['lien'], 0)] for q in AK.QA],
           [12, 20, 20, 26], {2: "#,##0", 3: "#,##0", 4: "#,##0"})
    D = wb.create_sheet("Dealer premium 2025")
    _sheet(D, "Retail contracts sold in 2025: premium received from dealers to 30 September 2026",
           ["Quarter of sale", "Premium owed (USD)", "Received (USD)", "Outstanding (USD)"],
           [[q, round(F['ask_b'][q]['owed'], 0), round(F['ask_b'][q]['received'], 0), round(F['ask_b'][q]['outstanding'], 0)]
            for q in AK.QB], [16, 20, 20, 20], {2: "#,##0", 3: "#,##0", 4: "#,##0"})
    N = wb.create_sheet("Notes")
    for line in (
            "Claims cost is payments on covered repairs net of the deductible and of reversed lines; goodwill is "
            "excluded (rate manual 4). Miles under cover run from the sale to each contract's last odometer reading "
            "in cover (CarLink closing report, repair orders, uploads; Fareline's weekly feed for its drivers), "
            "stopped at 75,000 miles; contract numbers before March 2024 joined through the crosswalk.",
            "Each 2026 Fareline contract stands for a 2027 sale: its odometer at sale and its annual mileage from "
            "the weekly feed, run to the earlier of 75,000 miles and 60 months through the cost per mile at each "
            "odometer band.",
            "Cancellations: the cancellation date is the first request received; a lienholder's payoff request is "
            "the contract on that VIN carrying that lienholder. Refunds per the terms (form DP-T 4.2, section 5).",
            "Dealer premium: a payment the bank returned (ACH returns report, matched on dealer, settlement month "
            "and amount) is not received; a bordereau re-sent after a return is.",
            "Prepared by the rideshare and fleet programme for the product committee, October 2026."):
        N.append([line])
    N.column_dimensions["A"].width = 130
    for row in N.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def deck(path, F):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt
    g = F['g']
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    ink, rust, grey = RGBColor(0x1D, 0x23, 0x30), RGBColor(0x7A, 0x4B, 0x1E), RGBColor(0x5A, 0x62, 0x70)

    def text(s, x, y, w, h, t, size, bold=False, color=ink):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = t
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        return tb

    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 0.6, 0.4, 12.0, 1.0, "Renew Fareline on 1 January 2027 with the maximum high-use surcharge", 28, True, rust)
    text(s, 0.6, 1.3, 12.0, 0.9,
         "A 2027 Fareline contract will cost about %d per cent of its price in claims at current rates and about %d "
         "per cent with the 40 per cent surcharge, against the rate manual's 70 per cent standard." % (r5(g['lr']), r5(g['lr_s'])),
         16, False, ink)
    rows = [("Election", g['call']), ("Lifetime loss ratio, current price", "about %d%%" % r5(g['lr'])),
            ("Lifetime loss ratio, surcharged price", "about %d%%" % r5(g['lr_s'])),
            ("Claims over a contract's cover (average)", "USD %s" % format(int(round(g['cost'], -2)), ",")),
            ("Contract price at current rates (average)", "USD %s" % format(int(round(g['price'], -1)), ","))]
    tb = s.shapes.add_table(len(rows), 2, Inches(0.6), Inches(2.5), Inches(8.5), Inches(2.6)).table
    tb.columns[0].width, tb.columns[1].width = Inches(4.6), Inches(3.9)
    tb.first_row = False
    tb.horz_banding = False
    for i, (a, b) in enumerate(rows):
        for j, v in enumerate((a, b)):
            c = tb.cell(i, j)
            c.text = v
            c.text_frame.paragraphs[0].font.size = Pt(14)
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor(0xF3, 0xEC, 0xE4) if i % 2 else RGBColor(0xFF, 0xFF, 0xFF)
            c.text_frame.paragraphs[0].font.color.rgb = ink
            c.text_frame.paragraphs[0].font.bold = (j == 1)
    text(s, 0.6, 5.5, 12.0, 1.2,
         "Fareline drivers buy cover on cars at 50,000 to 100,000 miles and drive about %s miles a year, so the "
         "75,000 covered miles run out in about %d months, at odometers where repairs come many times as often "
         "and cost more. The channel's own six months have seen only the start of that cover." % (
             format(int(round(g['rate'], -3)), ","), int(round(g['months']))), 13, False, grey)
    text(s, 0.6, 6.9, 12.0, 0.4, "Source: Copperline claims and odometer records to 30 September 2026; rate manual 2026, "
         "section 5.", 10, False, grey)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 0.6, 0.4, 12.0, 0.9, "Repairs cost per mile climbs past 70,000 miles, where a Fareline cover is spent", 24,
         True, rust)
    cd = CategoryChartData()
    cats, vals = [], []
    for j in range(len(g['cpm'])):
        lo = int(g['edges'][j])
        if g['miles'][j] <= 0 or lo >= 180000:
            continue
        cats.append("%d-%d" % (lo // 1000, lo // 1000 + 10))
        vals.append(round(float(g['cpm'][j]) * 1000, 1))
    cd.categories = cats
    cd.add_series("Claims cost per 1,000 miles (USD)", vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.4), Inches(12.0), Inches(4.6), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.plots[0].gap_width = 40
    ser = ch.plots[0].series[0]
    for k, cat in enumerate(cats):
        pt = ser.points[k]
        pt.format.fill.solid()
        lo = int(cat.split("-")[0]) * 1000
        pt.format.fill.fore_color.rgb = rust if 50000 <= lo < 175000 else RGBColor(0xC9, 0xB8, 0xA6)
    va = ch.value_axis
    va.has_title = True
    va.axis_title.text_frame.text = "Claims cost per 1,000 miles (USD)"
    va.axis_title.text_frame.paragraphs[0].font.size = Pt(12)
    va.tick_labels.font.size = Pt(11)
    ch.category_axis.tick_labels.font.size = Pt(11)
    ch.category_axis.has_title = True
    ch.category_axis.axis_title.text_frame.text = "Odometer band (thousand miles), whole book 2019 to September 2026"
    ch.category_axis.axis_title.text_frame.paragraphs[0].font.size = Pt(12)
    text(s, 0.6, 6.15, 12.0, 0.8,
         "Darker bars: the odometers a Fareline cover runs through (50,000 to 175,000 miles). Each 2026 Fareline "
         "contract is run from its odometer at sale at its own annual mileage to 75,000 miles or 60 months.", 12, False, grey)
    prs.save(path)
    _pin_embedded(path)


def _pin_embedded(path, stamp="2026-10-14T10:05:00Z"):
    """the chart's embedded workbook carries its own save time; pin it so two builds are byte-identical"""
    with zipfile.ZipFile(path) as z:
        items = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for name, data in items:
            if name.startswith("ppt/embeddings/") and name.endswith(".xlsx"):
                with zipfile.ZipFile(io.BytesIO(data)) as zi:
                    parts = [(i.filename, zi.read(i.filename)) for i in zi.infolist()]
                inner = io.BytesIO()
                with zipfile.ZipFile(inner, "w", zipfile.ZIP_DEFLATED) as zw:
                    for n2, d2 in parts:
                        if n2 == "docProps/core.xml":
                            d2 = re.sub(rb"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", stamp.encode(), d2)
                        zw.writestr(zipfile.ZipInfo(n2, date_time=(1980, 1, 1, 0, 0, 0)), d2)
                data = inner.getvalue()
            zo.writestr(zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0)), data)
    with open(path, "wb") as f:
        f.write(out.getvalue())
