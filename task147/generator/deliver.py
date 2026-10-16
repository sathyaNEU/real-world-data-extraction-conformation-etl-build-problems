"""The goldens, computed from the shipped pack: the executive housekeeper's block workbook (the block week and the next
week with their forecast room-nights above the crew's capacity, last August's services and overtime by roster week,
and the 2026 block's billable hours and charge by day) and the one-page note with the nightly forecast chart; and
generator/figures.json, every figure the write-up states."""
from __future__ import annotations

import datetime as dt
import io
import json
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt               # noqa: E402
from matplotlib import font_manager            # noqa: E402

import calc as C                               # noqa: E402
import common as K                             # noqa: E402
import postprocess as PP                       # noqa: E402

GOLD_BOOK = "housekeeping_block_2027.xlsx"
GOLD_NOTE = "housekeeping_block_2027.pdf"
for _f in ("Carlito-Regular.ttf", "Carlito-Bold.ttf"):
    font_manager.fontManager.addfont("/usr/share/fonts/truetype/crosextra/" + _f)
INK, INK2, MUTED, GRID, AXIS, PAPER = "#0b0b0b", "#52514e", "#898781", "#e1e0d9", "#c3c2b7", "#fcfcfb"
ACCENT, REST = "#2a78d6", "#a9a7a0"          # emphasis: the block week in slot 1, the other nights in gray


def short_week(w):
    m, s = C.WEEKS[w]
    return "%d to %d %s" % (m.day, s.day, s.strftime("%b")) if m.month == s.month else \
        "%d %s to %d %s" % (m.day, m.strftime("%b"), s.day, s.strftime("%b"))


def long_week(m, s):
    if m.month == s.month:
        return "Monday %d to Sunday %d %s" % (m.day, s.day, s.strftime("%B %Y"))
    return "Monday %d %s to Sunday %d %s" % (m.day, m.strftime("%B"), s.day, s.strftime("%B %Y"))


def compute(target):
    V = C.from_pack(target)
    L = C.ladder(V)
    cm = C.components(V, L)
    se = C.precision(V)
    a = C.ask_a(target)
    b = C.ask_b(target)
    return dict(V=V, L=L, cm=cm, se=se, a=a, b=b, b_week_charge=C.week_charge(b),
                pick=L["pick"], runner=L["runner"],
                pick_fig=K.round_to(L["gw"][L["pick"]], 100), runner_fig=K.round_to(L["gw"][L["runner"]], 100),
                default_pick=C.top(L["default"])[0], default_fig=K.round_to(max(L["default"]), 100),
                default_runner_fig=K.round_to(sorted(L["default"])[-2], 100),
                thru_w3=K.round_to(cm["thru_w3"], 10), nightly=[float(x) for x in L["g"]])


def n0(v):
    return "{:,}".format(int(v))


def hrs(v):
    return ("%.2f" % v)


def figures(c):
    L, cm = c["L"], c["cm"]
    weeks26 = K.roster_weeks(2026)
    return {
        "call": {"block_week": C.WEEK_NAMES[c["pick"]] + " 2027",
                 "room_nights_above_capacity": n0(c["pick_fig"]),
                 "next_week": C.WEEK_NAMES[c["runner"]] + " 2027", "next_week_room_nights": n0(c["runner_fig"])},
        "critical_components": {
            "full_nights": "Friday 13 and Saturday 14 August 2027",
            "rooms_on_the_book_full_nights": [n0(x) for x in cm["book_full"]],
            "concert_week_forecast": n0(c["runner_fig"]),
            "closing_week_forecast": n0(c["pick_fig"]),
        },
        "steps": {
            "extract_taken": "06:12 on Thursday 1 July 2027",
            "crew_capacity": "380", "hotel_rooms": "480",
            "past_seasons": "2022 to 2026",
            "comparable_point": "the Thursday five weeks before each year's opening night",
            "blocked_stays_room_nights_concert_week": n0(c["thru_w3"]),
        },
        "default_path": {"week": C.WEEK_NAMES[c["default_pick"]], "room_nights": n0(c["default_fig"]),
                         "its_runner_up_room_nights": n0(c["default_runner_fig"])},
        "precision": {"closing_week_standard_error": "%d" % round(c["se"][4]),
                      "concert_week_standard_error": "%d" % round(c["se"][2])},
        "season_window": "Monday 26 July to Sunday 29 August 2027",
        "ask_a_2026_services_and_overtime": [
            {"week": long_week(m, s), "crew_departure_cleans": n0(c["a"][w]["crew_dep"]),
             "crew_stayover_services": n0(c["a"][w]["crew_stay"]),
             "brennock_departure_cleans": n0(c["a"][w]["brn_dep"]),
             "brennock_stayover_services": n0(c["a"][w]["brn_stay"]),
             "crew_overtime_hours": hrs(c["a"][w]["ot"])} for w, (m, s) in enumerate(weeks26)],
        "ask_b_2026_block_billable_hours": [{"day": d.strftime("%A ") + "%d %s" % (d.day, d.strftime("%B %Y")),
                                             "hours": hrs(v[0])} for d, v in sorted(c["b"].items())],
        "ask_b_week": {"hours": hrs(sum(v[0] for v in c["b"].values())),
                       "charge_before_vat_eur": "{:,.2f}".format(c["b_week_charge"])},
    }


# ---------------------------------------------------------------------------------------------------- chart
def chart_png(c):
    plt.rcParams.update({"font.family": "Carlito", "font.size": 9.0, "axes.edgecolor": AXIS, "axes.labelcolor": INK2,
                         "xtick.color": INK2, "ytick.color": INK2})
    fig, ax = plt.subplots(figsize=(7.0, 3.05), dpi=180)
    fig.patch.set_facecolor(PAPER)
    ax.set_facecolor(PAPER)
    g = c["nightly"]
    xs = list(range(len(g)))
    pick = c["pick"]
    cols = [ACCENT if 7 * pick <= i < 7 * pick + 7 else REST for i in xs]
    ax.bar(xs, g, width=0.72, color=cols, zorder=3, linewidth=0)
    ax.axhline(K.CREW, color=INK, linewidth=1.1, zorder=4)
    ax.axhline(K.CAP, color=MUTED, linewidth=0.8, zorder=2)
    fi = c["L"]["parts"]["fi"]
    if fi:
        xm = sum(fi) / len(fi)
        ax.annotate("13 and 14 Aug full on the book", (xm, g[fi[0]] + 2), xytext=(xm - 7.0, 512), fontsize=8.0,
                    color=INK2, arrowprops=dict(arrowstyle="-", color=MUTED, linewidth=0.8), ha="center")
    xb = 7 * pick + 3
    ax.text(xb, 512, "Block week: about %s room-nights above 380" % n0(c["pick_fig"]), fontsize=8.4, color=INK,
            ha="center", va="bottom", fontweight="bold")
    for w in range(1, 5):
        ax.axvline(7 * w - 0.5, color=GRID, linewidth=0.8, zorder=1)
    ax.set_xticks([7 * w + 3 for w in range(5)])
    ax.set_xticklabels(["Roster week %d\n%s" % (w + 1, short_week(w)) for w in range(5)], fontsize=8.0)
    ax.set_xlim(-0.7, len(g) - 0.3)
    ax.set_ylim(0, 545)
    ax.set_yticks([0, 100, 200, 300, K.CREW, K.CAP])
    ax.set_yticklabels(["0", "100", "200", "300", "Crew's capacity 380", "Hotel's rooms 480"])
    for t in ax.get_yticklabels()[4:]:
        t.set_color(INK)
    ax.yaxis.grid(True, color=GRID, linewidth=0.7, zorder=0)
    ax.set_ylabel("Forecast occupied rooms a night", fontsize=8.4)
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.spines["bottom"].set_color(AXIS)
    ax.tick_params(length=0)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor=PAPER, metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


# ---------------------------------------------------------------------------------------------------- the note
def write_note(path, c, png):
    from reportlab.lib.units import cm
    from reportlab.platypus import Image, Paragraph, Spacer

    import docs as DOC
    from reportlab.lib.styles import ParagraphStyle
    pick, run = c["pick"], c["runner"]
    s = [Paragraph("The Weighhouse Hotel &middot; Housekeeping &middot; Festival season 2027", DOC.H1S),
         Paragraph("Brennock's block goes to the closing week, Monday 23 to Sunday 29 August" if pick == 4 else
                   "Brennock's block goes to %s" % C.WEEK_NAMES[pick], DOC.H1),
         Paragraph("Thursday 1 July 2027 &middot; on the book at 06:12 this morning &middot; Fergus Roddy, Executive "
                   "Housekeeper, for Edel McCay, Rooms Division Manager", DOC.H1S),
         Paragraph("We confirm the block to Brennock for <b>%s</b>. On this morning's book the week is forecast to "
                   "carry about <b>%s room-nights</b> above the crew's 380 rooms a day, the most of the five roster "
                   "weeks. The week that comes next is %s, at about %s." % (
                       C.WEEK_NAMES[pick], n0(c["pick_fig"]), C.WEEK_NAMES[run], n0(c["runner_fig"])), DOC.BODY),
         Paragraph("The concert week looks the biggest on a night-by-night pickup, which puts it at about %s, but "
                   "Friday 13 and Saturday 14 August are already full with the two stadium concerts. Many of that "
                   "week's late bookings in past Augusts were festival guests staying the week, and a stay is only "
                   "sold if every night of it has a room, so this year they cannot book through the concert nights: "
                   "in the last five seasons the stays booked after the confirmation date that ran through the same "
                   "Friday and Saturday carried about %s room-nights of the week's other five nights. Taking them "
                   "out leaves the concert week at about %s, its two full nights and little else. The closing "
                   "week has no full night and keeps its usual late bookings." % (
                       n0(c["default_fig"]), n0(c["thru_w3"]), n0(c["runner_fig"])), DOC.BODY),
         Spacer(1, 4),
         Image(io.BytesIO(png), width=16.6 * cm, height=16.6 * cm * 3.05 / 7.0),
         Paragraph("Forecast occupied rooms each night, Monday 26 July to Sunday 29 August 2027: the book at 06:12 on "
                   "1 July realised at each night's rate in 2022 to 2026, plus the stays booked after the confirmation "
                   "date in those seasons (aligned on the festival's opening), less those through the two full nights. "
                   "The full nights at 480 rooms less the event rate's no-shows.", DOC.SMALL),
         Spacer(1, 6),
         Paragraph("Room-nights above capacity are each night's occupied rooms above 380 summed over the week (HK-04 "
                   "3.2), rounded to the nearest hundred once. Last season's services and overtime and the 2026 block's "
                   "billing are in the workbook housekeeping_block_2027.xlsx.", DOC.SMALL)]
    doc = DOC._doc(path, "Brennock block 2027", K.PEOPLE["exec_hk"])
    f = DOC._footer("The Weighhouse Hotel  |  Housekeeping  |  Block confirmation note, 1 July 2027", "page ")
    doc.build(s, onFirstPage=f, onLaterPages=f)


# ---------------------------------------------------------------------------------------------------- the workbook
def write_book(path, c):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    head = PatternFill("solid", fgColor="EAE6E8")
    thin = Side(style="thin", color="A3A9B0")
    ws = wb.active
    ws.title = "Block week"
    ws["A1"] = "Brennock block, festival season 2027"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Forecast on the reservation book at 06:12 on Thursday 1 July 2027 (HK-04 4.1, 4.2)."
    ws.append([])
    ws.append(["", "Roster week", "Forecast room-nights above the crew's capacity"])
    for cl in ws[4]:
        cl.font = bold
        cl.fill = head
    ws.append(["Block week", C.WEEK_NAMES[c["pick"]] + " 2027", c["pick_fig"]])
    ws.append(["Next week", C.WEEK_NAMES[c["runner"]] + " 2027", c["runner_fig"]])
    ws["A5"].font = bold
    for r in (5, 6):
        ws.cell(row=r, column=3).number_format = "#,##0"
    ws.append([])
    ws.append(["Room-nights above capacity: each night's forecast occupied rooms above 380, summed over the roster "
               "week's seven nights (HK-04 3.2), to the nearest hundred."])
    ws.append(["Friday 13 and Saturday 14 August are full on the book (480 rooms); stays booked from today cannot "
               "run through them."])
    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 22
    ws["C4"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[4].height = 30

    ws2 = wb.create_sheet("2026 services")
    ws2["A1"] = "Last year's season by roster week: departure cleans and stayover services, and the crew's overtime"
    ws2["A1"].font = Font(bold=True, size=12)
    ws2["A2"] = ("Housekeeping app task log, 27/07/2026 to 30/08/2026, by the day the service was done (HK-04 2.4); "
                 "Brennock's temps as on Brennock's timesheets for the day; overtime by the date worked.")
    ws2.append([])
    hdr = ["Roster week", "Crew departure cleans", "Crew stayover services", "Brennock departure cleans",
           "Brennock stayover services", "Crew overtime hours"]
    ws2.append(hdr)
    for cl in ws2[4]:
        cl.font = bold
        cl.fill = head
        cl.alignment = Alignment(wrap_text=True, vertical="top")
    ws2.row_dimensions[4].height = 30
    tot = [0, 0, 0, 0, 0.0]
    for w, (m, s) in enumerate(K.roster_weeks(2026)):
        v = c["a"][w]
        row = [v["crew_dep"], v["crew_stay"], v["brn_dep"], v["brn_stay"], v["ot"]]
        tot = [a + b for a, b in zip(tot, row)]
        ws2.append(["%s to %s" % (m.strftime("%a %d/%m/%Y"), s.strftime("%a %d/%m/%Y"))] + row)
        r = ws2.max_row
        for col in range(2, 6):
            ws2.cell(row=r, column=col).number_format = "#,##0"
        ws2.cell(row=r, column=6).number_format = "0.00"
    ws2.append(["Season"] + tot)
    r = ws2.max_row
    for col in range(1, 7):
        ws2.cell(row=r, column=col).font = bold
        ws2.cell(row=r, column=col).border = Border(top=thin)
    for col in range(2, 6):
        ws2.cell(row=r, column=col).number_format = "#,##0"
    ws2.cell(row=r, column=6).number_format = "0.00"
    ws2.column_dimensions["A"].width = 31
    for col in "BCDEF":
        ws2.column_dimensions[col].width = 14
    ws2.freeze_panes = "B5"

    ws3 = wb.create_sheet("2026 block billing")
    ws3["A1"] = "Last year's block, 10 to 16 August 2026: the hours Brennock may bill for the block"
    ws3["A1"].font = Font(bold=True, size=12)
    ws3["A2"] = ("Framework agreement 4.1 to 4.3 on the staff access records (loading bay door LB-02 read on its UTC "
                 "clock); the week's charge at 5.1, before VAT.")
    ws3.append([])
    ws3.append(["Day", "Block room attendants", "Billable hours", "Rate (EUR an hour)", "Brennock invoice hours"])
    for cl in ws3[4]:
        cl.font = bold
        cl.fill = head
        cl.alignment = Alignment(wrap_text=True, vertical="top")
    ws3.row_dimensions[4].height = 30
    inv = c.get("invoice", {})
    th = 0.0
    for d, (h, ch) in sorted(c["b"].items()):
        n = c["block_n"][d]
        ws3.append([d, n, h, 28.20 if d.weekday() == 6 else 22.56, inv.get(d, (None,))[0]])
        r = ws3.max_row
        ws3.cell(row=r, column=1).number_format = "ddd dd/mm/yyyy"
        ws3.cell(row=r, column=3).number_format = "0.00"
        ws3.cell(row=r, column=4).number_format = "0.00"
        ws3.cell(row=r, column=5).number_format = "0.00"
        th += h
    ws3.append(["Week", None, th, None, sum(v[0] for v in inv.values()) if inv else None])
    r = ws3.max_row
    for col in range(1, 6):
        ws3.cell(row=r, column=col).font = bold
        ws3.cell(row=r, column=col).border = Border(top=thin)
    ws3.cell(row=r, column=3).number_format = "0.00"
    ws3.cell(row=r, column=5).number_format = "0.00"
    ws3.append([])
    ws3.append(["The week's charge before VAT (EUR)", None, c["b_week_charge"]])
    ws3.cell(row=ws3.max_row, column=1).font = bold
    ws3.cell(row=ws3.max_row, column=3).number_format = "#,##0.00"
    ws3.cell(row=ws3.max_row, column=3).font = bold
    ws3.column_dimensions["A"].width = 16
    for col in "BCDE":
        ws3.column_dimensions[col].width = 15
    wb.save(path)


def main(target, golden, figures_path, Wd=None):
    os.makedirs(golden, exist_ok=True)
    c = compute(target)
    ts = C.timesheets(target)
    c["block_n"] = {d: sum(1 for r in ts if r["kind"] == "Block" and r["date"] == d) for d in c["b"]}
    c["invoice"] = C.invoice_b(target)
    png = chart_png(c)
    write_note(os.path.join(golden, GOLD_NOTE), c, png)
    PP.normalize_pdf(os.path.join(golden, GOLD_NOTE), b"D:20270701102213+01'00'", b"D:20270701102213+01'00'")
    write_book(os.path.join(golden, GOLD_BOOK), c)
    PP.normalize_ooxml(os.path.join(golden, GOLD_BOOK), K.PEOPLE["exec_hk"], K.PEOPLE["exec_hk"],
                       "2027-07-01T08:05:31Z", "2027-07-01T09:19:47Z", excel=True)
    F = figures(c)
    with open(figures_path, "w", encoding="utf-8", newline="\n") as fh:
        json.dump(F, fh, indent=2, ensure_ascii=False)
        fh.write("\n")
    return c
