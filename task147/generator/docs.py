"""The pack's documents: the rooms division's housekeeping staffing standard HK-04, the framework agreement with
Brennock and Brennock's invoice for the 2026 block (PDF); the PMS code lists, the night audit statistics, the sales
events diary, the HR staff list, Brennock's timesheets and the competitor rate shop (xlsx). The standards' text is
kept here as data so that checks.py can quote it."""
from __future__ import annotations

import datetime as dt
import re

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import A4                                                # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import cm                                                    # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import cast as CA      # noqa: E402
import common as K     # noqa: E402
import records as R    # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/"
for _n, _f in (("Carlito", "crosextra/Carlito-Regular.ttf"), ("Carlito-Bold", "crosextra/Carlito-Bold.ttf"),
               ("Carlito-Italic", "crosextra/Carlito-Italic.ttf"),
               ("LibSerif", "liberation/LiberationSerif-Regular.ttf"),
               ("LibSerif-Bold", "liberation/LiberationSerif-Bold.ttf")):
    pdfmetrics.registerFont(TTFont(_n, FONT_DIR + _f))
SANS, SANS_B, SANS_I = "Carlito", "Carlito-Bold", "Carlito-Italic"
INK = colors.HexColor("#1d2329")
GREY = colors.HexColor("#5c636b")
RULE = colors.HexColor("#a3a9b0")
PALE = colors.HexColor("#ebeef0")
HOUSE = colors.HexColor("#5a2e3c")       # the hotel's plum
AGENCY = colors.HexColor("#1f5f74")      # Brennock's teal

BODY = ParagraphStyle("body", fontName=SANS, fontSize=10.0, leading=13.4, textColor=INK, spaceAfter=5)
H1 = ParagraphStyle("h1", fontName=SANS_B, fontSize=15.5, leading=19, textColor=HOUSE, spaceAfter=2)
H1S = ParagraphStyle("h1s", fontName=SANS, fontSize=9.2, leading=12.0, textColor=GREY, spaceAfter=10)
H2 = ParagraphStyle("h2", fontName=SANS_B, fontSize=11.4, leading=14.5, textColor=HOUSE, spaceBefore=8, spaceAfter=3)
H3 = ParagraphStyle("h3", fontName=SANS_B, fontSize=10.0, leading=13.0, textColor=INK, spaceBefore=4, spaceAfter=2)
CELL = ParagraphStyle("cell", fontName=SANS, fontSize=9.0, leading=11.0, textColor=INK)
CELLR = ParagraphStyle("cellr", parent=CELL, alignment=2)
CELLB = ParagraphStyle("cellb", fontName=SANS_B, fontSize=9.0, leading=11.0, textColor=INK)
CELLBR = ParagraphStyle("cellbr", parent=CELLB, alignment=2)
SMALL = ParagraphStyle("small", fontName=SANS, fontSize=8.4, leading=10.8, textColor=GREY, spaceBefore=3)
AG_BODY = ParagraphStyle("agb", fontName="LibSerif", fontSize=10.2, leading=13.6, textColor=INK, spaceAfter=5)
AG_H = ParagraphStyle("agh", fontName="LibSerif-Bold", fontSize=10.6, leading=14, textColor=INK, spaceBefore=7,
                      spaceAfter=2)
AG_T = ParagraphStyle("agt", fontName="LibSerif-Bold", fontSize=14, leading=18, textColor=INK, spaceAfter=4,
                      alignment=1)
AG_S = ParagraphStyle("ags", fontName="LibSerif", fontSize=10.2, leading=13.6, textColor=INK, spaceAfter=10,
                      alignment=1)


def _doc(path, title, author, margins=2.2):
    return SimpleDocTemplate(path, pagesize=A4, leftMargin=margins * cm, rightMargin=margins * cm, topMargin=1.9 * cm,
                             bottomMargin=2.0 * cm, title=title, author=author, subject="", creator="", producer="")


def _footer(left, right, font=SANS):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 7.8)
        canvas.setFillColor(GREY)
        canvas.drawString(2.2 * cm, 1.2 * cm, left)
        canvas.drawRightString(A4[0] - 2.2 * cm, 1.2 * cm, "%s%s" % (right, doc.page))
        canvas.restoreState()
    return draw


def _table(rows, widths, right=(), bold_last=False, head_bg=PALE):
    data = []
    for i, r in enumerate(rows):
        line = []
        for j, c in enumerate(r):
            hb = i == 0 or (bold_last and i == len(rows) - 1)
            st = (CELLBR if j in right else CELLB) if hb else (CELLR if j in right else CELL)
            line.append(Paragraph(str(c), st))
        data.append(line)
    t = Table(data, colWidths=widths, hAlign="LEFT")
    t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, 0), 0.6, RULE), ("LINEBELOW", (0, -1), (-1, -1), 0.6, RULE),
                           ("LINEABOVE", (0, 0), (-1, 0), 0.6, RULE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("BACKGROUND", (0, 0), (-1, 0), head_bg), ("TOPPADDING", (0, 0), (-1, -1), 2.2),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2)] +
                          ([("LINEABOVE", (0, -1), (-1, -1), 0.5, RULE)] if bold_last else [])))
    return t


def eur(v):
    return "{:,.2f}".format(v)


# ---------------------------------------------------------------------------------------------------- HK-04
F_HK04 = "HK-04_housekeeping_staffing_festival_season.pdf"
HK04 = [
    ("1", "Scope", None, [
        "This standard sets how housekeeping is staffed in the festival season and how Brennock's block of temporary "
        "room attendants is placed. It applies from the Monday of the week before the opening week of the Callowfort "
        "Festival of Arts to the festival's last Sunday."]),
    ("2", "Definitions", None, []),
    ("2.1", "Roster weeks", "weeks", [
        "The festival season is rostered in five roster weeks, each Monday to Sunday: the week before the opening "
        "week, the opening week (the week of the festival's opening night) and the three weeks after it. They are "
        "numbered 1 to 5 in every season."]),
    ("2.2", "Nights", "nights", [
        "A roster week covers the rooms occupied on its seven nights, Monday night to Sunday night."]),
    ("2.3", "Occupied rooms", "occupied", [
        "An occupied room is a room with a guest in it overnight. Complimentary and house-use rooms are occupied "
        "rooms. No-show and day-use rooms are not."]),
    ("2.4", "Services", "services", [
        "A departure clean is the full clean of a room after its guest checks out. A stayover service is the daily "
        "service of an occupied room whose guest stays on. A room is not serviced on a day its guest asks not to be "
        "disturbed or declines the service. A re-clean after a failed inspection is part of the service it repeats. "
        "Services are counted on the day they are done."]),
    ("2.5", "Crew", "crew", [
        "The crew are the hotel's own room attendants. Brennock Hospitality Staffing supplies temporary room "
        "attendants under the framework agreement of March 2025."]),
    ("3", "Capacity", None, []),
    ("3.1", "The crew's capacity", "capacity", ["The crew can clean 380 occupied rooms a day."]),
    ("3.2", "Rooms above capacity", "above", [
        "A night's rooms above capacity are its occupied rooms above 380. A roster week's rooms above capacity are "
        "the sum of its seven nights' rooms above capacity, in room-nights."]),
    ("4", "The agency block", None, []),
    ("4.1", "Placing the block", "block", [
        "Brennock supplies one block of temporary room attendants for one roster week of the season. The block goes "
        "to the roster week with the most room-nights above capacity, forecast for the season."]),
    ("4.2", "Confirmation", "confirm", [
        "The block is confirmed to Brennock on the Thursday five weeks before the festival's opening night, on the "
        "book as it stands that morning."]),
    ("4.3", "Size", "size", [
        "The block's size and daily shifts are agreed with Brennock when the week is confirmed."]),
    ("5", "Overtime", "overtime", ["Room attendants' overtime is approved by the executive housekeeper."]),
]


def hk04_sentences(key):
    for _, _, k, paras in HK04:
        if k == key:
            return [s for p in paras for s in re.split(r"(?<=\.)\s+(?=[A-Z])", p)]
    raise KeyError(key)


def hk04_text():
    return " ".join(" ".join(p) for _, _, _, p in HK04)


def write_hk04(path):
    s = [Paragraph("The Weighhouse Hotel &middot; Rooms Division &middot; Standard HK-04, issue 3", H1S),
         Paragraph("Housekeeping staffing for the festival season", H1),
         Paragraph("Effective 1 April 2025. Owner: Edel McCay, Rooms Division Manager. Replaces issue 2 (February "
                   "2022).", H1S)]
    for num, head, _, paras in HK04:
        s.append(Paragraph("%s&nbsp;&nbsp;%s" % (num, head), H2 if "." not in num else H3))
        for p in paras:
            s.append(Paragraph(p.replace("&", "&amp;"), BODY))
    doc = _doc(path, "HK-04 Housekeeping staffing for the festival season", K.PEOPLE["rooms_div"])
    f = _footer("The Weighhouse Hotel  |  HK-04 issue 3  |  Controlled copy: Rooms Division shared drive", "page ")
    doc.build(s, onFirstPage=f, onLaterPages=f)


# ---------------------------------------------------------------------------------------------------- agreement
F_AGREEMENT = "brennock_framework_agreement_2025.pdf"
AGREEMENT = [
    ("1", "Term", None, ["This agreement runs from 1 April 2025 to 31 December 2027 and replaces the agreement of "
                          "May 2022."]),
    ("2", "Bookings", None, [
        "2.1 The hotel books temporary room attendants by written order: regular placements for the season, day "
        "bookings, and one block booking for one week of the festival season.",
        "2.2 Every booking sets its booked shifts: the date, the booked start and the booked end of each shift.",
        "2.3 A room attendant without a staff badge signs in at the hotel's reception for a day badge."]),
    ("3", "Timesheets", None, [
        "3.1 Brennock keeps a timesheet of every booked shift, signed off by the hotel's supervisor on duty.",
        "3.2 Timesheets record the booked shifts. The record of attendance is the hotel's staff access system."]),
    ("4", "Billable hours", "billable", [
        "4.1 The billable time of a booked shift runs from the later of the booked start and the room attendant's "
        "first entry that day to the earlier of the booked end and the room attendant's last exit that day, and is "
        "counted in whole quarter hours, rounded down.",
        "4.2 Thirty minutes' unpaid break is deducted from a shift of more than six hours of billable time.",
        "4.3 A booked shift that the room attendant attends is billed at no less than four hours.",
        "4.4 Time outside a booked shift is not billable unless the hotel has booked it in writing."]),
    ("5", "Charges", "charges", [
        "5.1 Billable hours are charged at EUR 22.56 an hour from 1 January 2026 (EUR 21.92 in 2025). Hours on "
        "Sundays and public holidays are charged at one and a quarter times the hourly charge.",
        "5.2 Charges exclude VAT, which is added at the standard rate.",
        "5.3 Brennock invoices each week's shifts by the Wednesday after the week, by booking."]),
    ("6", "Queries", None, [
        "6.1 The hotel may withhold the part of an invoice that bills more than this agreement allows and pays the "
        "balance within the agreed terms.",
        "6.2 Payment is due 30 days from the invoice date."]),
]


def agreement_sentences(key):
    for _, _, k, paras in AGREEMENT:
        if k == key:
            return list(paras)
    raise KeyError(key)


def write_agreement(path):
    s = [Paragraph("Framework agreement for the supply of temporary housekeeping staff", AG_T),
         Paragraph("between The Weighhouse Hotel Limited, Weighhouse Square, Callowfort (the hotel)<br/>and Brennock "
                   "Hospitality Staffing Limited, Unit 4, Tannery Lane, Callowfort (Brennock)", AG_S)]
    for num, head, _, paras in AGREEMENT:
        s.append(Paragraph("%s. %s" % (num, head), AG_H))
        for p in paras:
            s.append(Paragraph(p.replace("&", "&amp;"), AG_BODY))
    s.append(Spacer(1, 16))
    s.append(Paragraph("Signed for the hotel: Edel McCay, Rooms Division Manager. Signed for Brennock: Naoise Meskill, "
                       "Account Manager. 14 March 2025.", AG_BODY))
    doc = _doc(path, "Framework agreement, temporary housekeeping staff", K.PEOPLE["agency"], margins=2.6)
    f = _footer("Brennock Hospitality Staffing Ltd  |  Framework agreement WH-2025-03", "Page ", font="LibSerif")
    doc.build(s, onFirstPage=f, onLaterPages=f)


# ---------------------------------------------------------------------------------------------------- invoice
F_INVOICE = "brennock_invoice_BHS-INV-2026-0833.pdf"


def write_invoice(path, inv):
    s = [Paragraph("Brennock Hospitality Staffing Ltd", ParagraphStyle("bn", parent=H1, textColor=AGENCY)),
         Paragraph("Unit 4, Tannery Lane, Callowfort &middot; VAT no. IE3471185NH", H1S),
         Paragraph("<b>Invoice %s</b>" % inv["no"], ParagraphStyle("ib", parent=BODY, fontSize=12, leading=15)),
         Paragraph("Date: %s<br/>To: The Weighhouse Hotel Ltd, Accounts Payable, Weighhouse Square, Callowfort<br/>"
                   "Your order: WH-PO-26-0417<br/>Booking: block booking, room attendants, Monday 10 August to Sunday 16 "
                   "August 2026, booked shifts 08:00 to 16:00" % inv["date"].strftime("%d/%m/%Y"), BODY),
         Spacer(1, 6)]
    rows = [["Date", "Room attendants", "Hours", "Rate (EUR)", "Amount (EUR)"]]
    for d, n, hrs, rate, amt in inv["lines"]:
        rows.append([d.strftime("%a %d/%m/%Y"), str(n), "%.2f" % hrs, "%.2f" % rate, eur(amt)])
    rows.append(["", "", "", "Net", eur(inv["net"])])
    rows.append(["", "", "", "VAT 23%", eur(inv["vat"])])
    rows.append(["", "", "", "Total", eur(inv["total"])])
    s.append(_table(rows, [4.2 * cm, 3.0 * cm, 2.4 * cm, 2.6 * cm, 3.2 * cm], right=(1, 2, 3, 4), bold_last=True,
                    head_bg=colors.HexColor("#e2eef1")))
    s.append(Spacer(1, 8))
    s.append(Paragraph("Hours: booked shifts of 8.00 hours less 0.50 hours' break. Sunday at 1.25 times the rate. "
                       "Payment 30 days from the invoice date by bank transfer to the account on your statement.", SMALL))
    doc = _doc(path, "Invoice %s" % inv["no"], "Brennock Hospitality Staffing Ltd")
    f = _footer("Brennock Hospitality Staffing Ltd, registered in Ireland no. 612904", "Page ")
    doc.build(s, onFirstPage=f, onLaterPages=f)


# ---------------------------------------------------------------------------------------------------- workbooks
def _wb():
    import openpyxl
    return openpyxl.Workbook()


def _widths(ws, widths):
    import openpyxl
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w


def write_codes(path):
    from openpyxl.styles import Font
    wb = _wb()
    ws = wb.active
    ws.title = "Rate codes"
    ws.append(["Rate code", "Description", "Cancellation", "Deposit", "Default market"])
    for r in R.RATE_CODES:
        ws.append(list(r))
    _widths(ws, [10, 44, 58, 32, 15])
    ws2 = wb.create_sheet("Channel codes")
    ws2.append(["Channel code", "Description", "Valid from", "Valid to", "Replaced by"])
    for c in R.CHANNELS:
        ws2.append([c[0], c[1], dt.date.fromisoformat(c[2]), dt.date.fromisoformat(c[3]) if c[3] else None,
                    c[4] or None])
        for col in (3, 4):
            ws2.cell(row=ws2.max_row, column=col).number_format = "dd/mm/yyyy"
    _widths(ws2, [13, 58, 12, 12, 12])
    ws3 = wb.create_sheet("Market segments")
    ws3.append(["Market code", "Description"])
    for m in R.MARKETS:
        ws3.append(list(m))
    _widths(ws3, [12, 30])
    for w in (ws, ws2, ws3):
        for c in w[1]:
            c.font = Font(bold=True)
        w.freeze_panes = "A2"
    wb.save(path)


AUD_OLD = ["Date", "Day", "Rooms Sold", "Comp", "House Use", "Total Occ", "Occ %", "Day Use", "Arrivals", "Departures",
           "No Shows", "Room Revenue", "ADR", "RevPAR"]
AUD_NEW = ["Business date", "Rooms available", "Occupied rooms", "of which complimentary", "of which house use",
           "Day use", "Arrival rooms", "Departure rooms", "No-show rooms", "Occupancy %", "Room revenue (EUR)",
           "ADR (EUR)", "RevPAR (EUR)"]


def write_audit(path, rows):
    from openpyxl.styles import Alignment, Font
    wb = _wb()
    first = True
    for y in range(2022, 2028):
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = str(y)
        yr = [r for r in rows if r["date"].year == y]
        if y <= 2023:
            ws.append(["THE WEIGHHOUSE HOTEL - NIGHT AUDIT STATISTICS %d" % y])
            ws["A1"].font = Font(bold=True)
            ws.append(AUD_OLD)
            for c in ws[2]:
                c.font = Font(bold=True)
            for r in yr:
                sold = r["occ"] - r["comp"] - r["hse"]
                ws.append([r["date"], r["date"].strftime("%a").upper(), sold, r["comp"], r["hse"], r["occ"],
                           r["occ"] / 480.0, r["dayuse"], r["arr"], r["dep"], r["ns"], r["rev"],
                           r["rev"] / sold if sold else 0.0, r["rev"] / 480.0])
                n = ws.max_row
                ws.cell(row=n, column=1).number_format = "dd/mm/yyyy"
                ws.cell(row=n, column=7).number_format = "0.0%"
                for col in (12, 13, 14):
                    ws.cell(row=n, column=col).number_format = "#,##0.00"
            ws.freeze_panes = "A3"
            _widths(ws, [11, 6, 11, 7, 10, 10, 8, 9, 9, 11, 10, 14, 9, 9])
        else:
            ws.append(["The Weighhouse Hotel", None, None, "Night audit statistics %d" % y])
            ws.append(["Occupied rooms include complimentary and house-use rooms. ADR on paying occupied rooms. "
                       "Revenue VAT included. Closed by the night audit each business date."])
            ws.append([])
            ws.append(AUD_NEW)
            for c in ws[4]:
                c.font = Font(bold=True)
                c.alignment = Alignment(wrap_text=True, vertical="top")
            for r in yr:
                paid = r["occ"] - r["comp"] - r["hse"]
                ws.append([r["date"], 480, r["occ"], r["comp"], r["hse"], r["dayuse"], r["arr"], r["dep"], r["ns"],
                           r["occ"] / 480.0, r["rev"], round(r["rev"] / paid, 2) if paid else 0.0,
                           round(r["rev"] / 480.0, 2)])
                n = ws.max_row
                ws.cell(row=n, column=1).number_format = "dd/mm/yyyy"
                ws.cell(row=n, column=10).number_format = "0.00%"
                for col in (11, 12, 13):
                    ws.cell(row=n, column=col).number_format = "#,##0.00"
            ws.freeze_panes = "B5"
            _widths(ws, [13, 10, 10, 13, 11, 9, 10, 11, 10, 11, 14, 10, 11])
    wb.save(path)


def write_diary(path, rows):
    from openpyxl.styles import Font
    wb = _wb()
    ws = wb.active
    ws.title = "Diary"
    ws.append(["Sales office events diary"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["City events the sales office tracks for pricing and group space. Kept by reservations."])
    ws.append([])
    ws.append(["From", "To", "Event", "Venue", "Category", "Announced", "Notes"])
    for c in ws[4]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(list(r))
        n = ws.max_row
        for col in (1, 2, 6):
            ws.cell(row=n, column=col).number_format = "dd/mm/yyyy"
    ws.freeze_panes = "A5"
    _widths(ws, [11, 11, 36, 27, 12, 11, 40])
    wb.save(path)


def write_hr(path, rows):
    from openpyxl.styles import Font
    wb = _wb()
    ws = wb.active
    ws.title = "Housekeeping"
    ws.append(["Housekeeping department: staff list as at 28/06/2027"])
    ws["A1"].font = Font(bold=True)
    ws.append(["Source: HR system, current employees. Prepared by Grainne McGuinn, HR and Payroll."])
    ws.append([])
    ws.append(["Employee no", "Surname", "First name", "Job title", "Department", "Contract", "Start date",
               "Contracted hours"])
    for c in ws[4]:
        c.font = Font(bold=True)
    for emp, name, title, contract, start, hrs in rows:
        f, s = name.split(" ", 1)
        ws.append([emp, s, f, title, "Housekeeping", contract, start, hrs])
        ws.cell(row=ws.max_row, column=7).number_format = "dd/mm/yyyy"
    ws.freeze_panes = "A5"
    _widths(ws, [12, 16, 13, 22, 14, 22, 11, 10])
    wb.save(path)


def write_timesheets(path, rows):
    from openpyxl.styles import Font
    wb = _wb()
    weeks = sorted({r["week_end"] for r in rows})
    first = True
    for we in weeks:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "WE %s" % we.strftime("%d.%m.%y")
        ws.append(["Brennock Hospitality Staffing - timesheet summary"])
        ws["A1"].font = Font(bold=True)
        ws.append(["Client: The Weighhouse Hotel, Callowfort", None, None, "Week ending:", we])
        ws.cell(row=2, column=5).number_format = "dd/mm/yyyy"
        ws.append([])
        ws.append(["Worker ID", "Worker", "Booking", "Date", "Booked start", "Booked end", "Break (min)", "Hours",
                   "Signed off by", "Notes"])
        for c in ws[4]:
            c.font = Font(bold=True)
        for r in sorted([x for x in rows if x["week_end"] == we], key=lambda x: (x["date"], x["kind"], x["name"])):
            ws.append([r["bhs"], r["name"], r["kind"], r["date"], r["start"], r["end"], r["brk"], r["hours"],
                       r["approver"], r["note"] or None])
            ws.cell(row=ws.max_row, column=4).number_format = "dd/mm/yyyy"
            ws.cell(row=ws.max_row, column=8).number_format = "0.00"
        ws.freeze_panes = "A5"
        _widths(ws, [11, 21, 9, 11, 11, 10, 10, 7, 15, 24])
    wb.save(path)


def write_rateshop(path, rows, hotels):
    from openpyxl.styles import Font
    wb = _wb()
    ws = wb.active
    ws.title = "Lowest rates"
    ws.append(["Rate shop: lowest public rate, 1 night, 2 adults, room only, EUR"])
    ws["A1"].font = Font(bold=True)
    ws.append(["Shopped 30/06/2027 07:00 on the hotels' own websites. Arrivals 26/07/2027 to 29/08/2027."])
    ws.append([])
    ws.append(["Arrival", "Day"] + hotels)
    for c in ws[4]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r[0], r[0].strftime("%a")] + r[1:])
        ws.cell(row=ws.max_row, column=1).number_format = "dd/mm/yyyy"
    ws.freeze_panes = "C5"
    _widths(ws, [11, 6] + [20] * len(hotels))
    wb.save(path)
