"""The independent verifier: reads only the bundle (target/ and golden/) with the standard library and openpyxl, on
its own code path (no generator module), and recomputes the call, the next week, the default path, and every Ask A
and Ask B figure, then compares them with the goldens and generator/figures.json."""
from __future__ import annotations

import csv
import datetime as dt
import json
import math
import os
import re
import sys
from collections import defaultdict

DATE0 = dt.date(2021, 1, 1)
CAP, CREW = 480, 380


def day(s):
    return (dt.date.fromisoformat(s) - DATE0).days


def first_friday(y):
    d = dt.date(y, 8, 1)
    while d.weekday() != 4:
        d += dt.timedelta(days=1)
    return d


def last_sunday(y, m):
    d = dt.date(y, m + 1, 1) - dt.timedelta(days=1)
    while d.weekday() != 6:
        d -= dt.timedelta(days=1)
    return d


def to_utc(local):
    """Irish summer time: an hour ahead of UTC from the last Sunday of March to the last Sunday of October."""
    y = local.year
    on = dt.datetime.combine(last_sunday(y, 3), dt.time(1)) <= local < dt.datetime.combine(last_sunday(y, 10), dt.time(2))
    return local - dt.timedelta(hours=1 if on else 0)


def cut(y):
    """The Thursday five weeks before the festival's opening night (the Thursday before the first Friday), at the
    time the 2027 extract was taken (06:12:44), in UTC as the extract writes it."""
    opening_night = first_friday(y) - dt.timedelta(days=1)
    return to_utc(dt.datetime.combine(opening_night - dt.timedelta(days=35), dt.time(6, 12, 44))).strftime(
        "%Y-%m-%d %H:%M:%S")


def season(y):
    f = first_friday(y)
    mon = f - dt.timedelta(days=f.weekday() + 7)
    return [mon + dt.timedelta(days=i) for i in range(35)]


def rnd(x, step):
    q = x / step
    return int(math.floor(q + 0.5)) * step


def main(target, golden, figures_path=None):
    rows = []
    with open(os.path.join(target, "pms_reservations_extract_2027-07-01.csv"), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            rows.append((int(r["confirmation_no"]), r["created_utc"], r["cancelled_utc"], day(r["arrival"]),
                         day(r["departure"]), r["status"], r["rate_code"]))
    changes = defaultdict(list)
    with open(os.path.join(target, "pms_amendment_log_2027-07-01.csv"), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            if r["field"] in ("ARRIVAL", "DEPARTURE"):
                changes[int(r["confirmation_no"])].append((r["changed_utc"], r["field"], day(r["old_value"])))
    for v in changes.values():
        v.sort()

    def tables(y, full):
        """Each aligned night of year y: the book at the cut (dates as they stood), the rooms of it that stayed on the
        night, the finals, the stays booked after the cut and those through a full night."""
        t = cut(y)
        nights = season(y)
        n0 = (nights[0] - DATE0).days
        z = lambda: [0.0] * 35
        book, real, final, new, thru = z(), z(), z(), z(), z()
        for conf, cr, cx, a, d, st, rc in rows:
            if d <= a:
                continue
            a_t, d_t = a, d
            for ct, fld, old in changes.get(conf, ()):
                if ct > t:
                    if fld == "ARRIVAL" and a_t == a:
                        a_t = old
                    if fld == "DEPARTURE" and d_t == d:
                        d_t = old
            stayed = st in ("OUT", "INH")
            alive = cr < t and (st != "CXL" or cx > t)
            if alive:
                for n in range(max(a_t, n0), min(d_t, n0 + 35)):
                    book[n - n0] += 1
                    if stayed and a <= n < d:
                        real[n - n0] += 1
            if stayed:
                through = any(a <= f < d for f in full)
                for n in range(max(a, n0), min(d, n0 + 35)):
                    final[n - n0] += 1
                    if cr >= t:
                        new[n - n0] += 1
                        if through:
                            thru[n - n0] += 1
        return dict(book=book, real=real, final=final, new=new, thru=thru)

    past = {}
    for y in range(2022, 2027):
        fo = (first_friday(y) - DATE0).days
        past[y] = tables(y, (fo + 7, fo + 8))
    now = tables(2027, ())
    book = now["book"]
    full = [i for i in range(35) if book[i] >= CAP]
    evt = [r for r in rows if r[6] == "EVT" and r[5] in ("OUT", "NS")]
    ns_share = sum(1 for r in evt if r[5] == "NS") / len(evt)
    Y = list(past)

    def weeks(v):
        return [sum(max(0.0, x - CREW) for x in v[7 * w:7 * w + 7]) for w in range(5)]

    gold, dflt = [], []
    for i in range(35):
        realised = sum(past[y]["real"][i] for y in Y) / sum(past[y]["book"][i] for y in Y)
        g = book[i] * realised + sum(past[y]["new"][i] - past[y]["thru"][i] for y in Y) / len(Y)
        if i in full:
            g = CAP * (1 - ns_share)
        gold.append(min(CAP, g))
        dflt.append(min(CAP, book[i] + sum(past[y]["final"][i] - past[y]["book"][i] for y in Y) / len(Y)))
    gw, dw = weeks(gold), weeks(dflt)
    order = sorted(range(5), key=lambda w: -gw[w])
    dorder = sorted(range(5), key=lambda w: -dw[w])
    out = dict(full=[season(2027)[i].isoformat() for i in full], gold=gw, default=dw, pick=order[0], next=order[1],
               pick_fig=rnd(gw[order[0]], 100), next_fig=rnd(gw[order[1]], 100), default_pick=dorder[0],
               default_fig=rnd(dw[dorder[0]], 100))

    # ---------------------------------------------------------------- Ask A
    users = {r["login"]: r for r in csv.DictReader(open(os.path.join(target, "hk_app_users_2027-06-28.csv"),
                                                         encoding="utf-8"))}
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(target, "brennock_timesheets_weighhouse_summer_2026.xlsx"), read_only=True)
    sheet_rows = []
    for ws in wb.worksheets:
        for r in ws.iter_rows(min_row=5, values_only=True):
            if r[0]:
                sheet_rows.append(r)
    agency_days = {(r[1], r[3].date().isoformat()) for r in sheet_rows}
    people = {u["display_name"].split(" ")[0]: u["display_name"] for u in users.values()
              if not u["login"].startswith("block")}
    for r in sheet_rows:
        people.setdefault(r[1].split(" ")[0], r[1])
    wk26 = season(2026)
    ask_a = [dict(crew_dep=0, crew_stay=0, brn_dep=0, brn_stay=0, ot=0.0) for _ in range(5)]
    with open(os.path.join(target, "hk_task_log_2026-07-27_2026-08-30.csv"), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            if r["status"] != "Completed" or r["task_type"] not in ("Departure", "Stayover"):
                continue
            d = dt.date.fromisoformat(r["service_date"])
            w = (d - wk26[0]).days // 7
            if not 0 <= w < 5:
                continue
            u = users[r["user_login"]]
            who = u["display_name"]
            if u["user_group"] == "Housekeeping Supervisors":
                named = [people[f] for f in people if re.search(r"\b%s\b" % re.escape(f), r["note"])]
                who = named[0]
            agency = (who, r["service_date"]) in agency_days or (u["user_group"] == "Agency - Brennock" and
                                                                  who == u["display_name"])
            ask_a[w]["%s_%s" % ("brn" if agency else "crew", "dep" if r["task_type"] == "Departure" else "stay")] += 1
    with open(os.path.join(target, "payroll_overtime_lines_2026-P14-P19.csv"), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            dd, mm, yy = r["reference"].split(" ")[1].split("/")
            d = dt.date(2000 + int(yy), int(mm), int(dd))
            w = (d - wk26[0]).days // 7
            if 0 <= w < 5:
                ask_a[w]["ot"] += float(r["hours"])
    out["ask_a"] = ask_a

    # ---------------------------------------------------------------- Ask B
    block = [r for r in sheet_rows if r[2] == "Block"]
    badge = {}
    with open(os.path.join(target, "reception_day_badge_log_2026-08.csv"), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            badge[(r["date"], r["issued_to"])] = r["badge"]
    events = defaultdict(list)
    with open(os.path.join(target, "staff_access_events_2026-08.csv"), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            t = dt.datetime.fromisoformat(r["event_time"])
            if r["reader"] == "LB-02":      # the loading bay's controller keeps UTC: an hour behind in August
                t += dt.timedelta(hours=1)
            events[(t.date().isoformat(), r["badge"])].append((t, r["direction"]))
    ask_b = {}
    for d in sorted({r[3].date() for r in block}):
        hours = 0.0
        for r in [x for x in block if x[3].date() == d]:
            ev = events[(d.isoformat(), badge[(d.isoformat(), r[1])])]
            first_in = min(t for t, k in ev if k == "IN")
            last_out = max(t for t, k in ev if k == "OUT")
            s = max(first_in, dt.datetime.combine(d, dt.time.fromisoformat(r[4])))
            e = min(last_out, dt.datetime.combine(d, dt.time.fromisoformat(r[5])))
            q = int(max(0.0, (e - s).total_seconds()) // 900) / 4.0
            if q > 6:
                q -= 0.5
            hours += max(q, 4.0)
        rate = 22.56 * (1.25 if d.weekday() == 6 else 1.0)
        ask_b[d.isoformat()] = (hours, math.floor(hours * rate * 100 + 0.5) / 100)
    out["ask_b"] = ask_b

    # ---------------------------------------------------------------- compare with the goldens and figures.json
    fails = []
    gb = openpyxl.load_workbook(os.path.join(golden, "housekeeping_block_2027.xlsx"), data_only=True)
    ws = gb["Block week"]
    names = ["Monday 26 July to Sunday 1 August", "Monday 2 to Sunday 8 August", "Monday 9 to Sunday 15 August",
             "Monday 16 to Sunday 22 August", "Monday 23 to Sunday 29 August"]
    if ws["B5"].value != names[out["pick"]] + " 2027" or ws["C5"].value != out["pick_fig"]:
        fails.append(("block week", ws["B5"].value, ws["C5"].value, names[out["pick"]], out["pick_fig"]))
    if ws["B6"].value != names[out["next"]] + " 2027" or ws["C6"].value != out["next_fig"]:
        fails.append(("next week", ws["B6"].value, ws["C6"].value, names[out["next"]], out["next_fig"]))
    ws2 = gb["2026 services"]
    for w in range(5):
        got = [ws2.cell(row=5 + w, column=c).value for c in range(2, 7)]
        want = [ask_a[w]["crew_dep"], ask_a[w]["crew_stay"], ask_a[w]["brn_dep"], ask_a[w]["brn_stay"], ask_a[w]["ot"]]
        if got[:4] != want[:4] or abs(got[4] - want[4]) > 1e-9:
            fails.append(("services week %d" % (w + 1), got, want))
    ws3 = gb["2026 block billing"]
    for i, (d, (h, ch)) in enumerate(sorted(ask_b.items())):
        got = ws3.cell(row=5 + i, column=3).value
        if abs(got - h) > 1e-9:
            fails.append(("block day %s" % d, got, h))
    week = math.floor(sum(h * 22.56 * (1.25 if dt.date.fromisoformat(d).weekday() == 6 else 1.0)
                          for d, (h, _) in ask_b.items()) * 100 + 0.5) / 100
    got = [r[2] for r in ws3.iter_rows(values_only=True) if r[0] == "The week's charge before VAT (EUR)"]
    if not got or abs(got[0] - week) > 0.005:
        fails.append(("week charge", got, week))
    out["week_charge"] = week
    if figures_path and os.path.exists(figures_path):
        F = json.load(open(figures_path, encoding="utf-8"))
        if F["call"]["room_nights_above_capacity"] != "{:,}".format(out["pick_fig"]) or \
                F["call"]["next_week_room_nights"] != "{:,}".format(out["next_fig"]) or \
                F["default_path"]["room_nights"] != "{:,}".format(out["default_fig"]):
            fails.append(("figures.json call", F["call"], out["pick_fig"], out["next_fig"]))
    if fails:
        for f in fails:
            print("VERIFY FAIL", f)
        raise SystemExit(1)
    print("verify.py: the call W%d (%.0f, about %d), next W%d (%.0f, about %d), the default W%d (%.0f, about %d); "
          "full nights %s; Ask A 25 figures and Ask B 8 figures recomputed and equal in the workbook" % (
              out["pick"] + 1, gw[out["pick"]], out["pick_fig"], out["next"] + 1, gw[out["next"]], out["next_fig"],
              out["default_pick"] + 1, dw[out["default_pick"]], out["default_fig"], ", ".join(out["full"])))
    return out


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    task = os.path.dirname(here)
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(task, "target"),
         sys.argv[2] if len(sys.argv) > 2 else os.path.join(task, "golden"),
         os.path.join(here, "figures.json"))
