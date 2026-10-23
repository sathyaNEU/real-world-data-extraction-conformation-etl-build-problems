"""task176 independent verifier (python3 verify.py <target>): reads only the shipped bundle, on its own code path
(pandas frames and its own Poisson fit by iteratively reweighted least squares, no import of calc.py), and
recomputes the call, its components and every ask figure, then compares them with generator/figures.json and the
golden workbook. It reads back every format the pack ships. Exit 1 on any mismatch."""
from __future__ import annotations

import datetime as dt
import json
import os
import sys

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
BAD = []


def say(label, cond, detail=""):
    print("%-4s %-50s %s" % ("PASS" if cond else "FAIL", label, detail))
    if not cond:
        BAD.append(label)


def parse_birth(s):
    s = s.strip()
    if "/" in s:
        m, d, y = s.split("/")
        return dt.date(int(y), int(m), int(d))
    return dt.date.fromisoformat(s)


def kyear(b):
    return b.year + 5 if (b.month, b.day) <= (9, 1) else b.year + 6


def mindex(d):
    return d.year * 12 + d.month - 1


def irls_poisson(X, y, expo, it=60):
    """Poisson regression with log exposure offset by iteratively reweighted least squares"""
    beta = np.zeros(X.shape[1])
    beta[0] = np.log(max(y.sum() / expo.sum(), 1e-9))
    for _ in range(it):
        eta = X @ beta + np.log(expo)
        mu = np.exp(eta)
        z = eta - np.log(expo) + (y - mu) / mu
        W = mu
        XtW = X.T * W
        new = np.linalg.solve(XtW @ X + 1e-10 * np.eye(X.shape[1]), XtW @ z)
        if np.max(np.abs(new - beta)) < 1e-10:
            beta = new
            break
        beta = new
    return beta


def main(target):
    task = os.path.dirname(target)
    figs = json.load(open(os.path.join(HERE, "figures.json")))
    lic = pd.read_csv(os.path.join(target, "fcc_licence_extract_20261231.csv"), dtype=str, keep_default_na=False)
    xw = pd.read_csv(os.path.join(target, "licence_number_crosswalk_2019.csv"), dtype=str)
    ws = pd.read_csv(os.path.join(target, "capacity_worksheet_children_2016_2026.csv"), dtype=str)
    m = dict(zip(xw.legacy_licence_no, xw.licence_no))
    ws["no"] = ws.licence_ref.map(lambda x: m.get(x, x))
    ws["rel"] = ws.relationship.str.strip().str.upper().map({"C": "OWN", "OWN CHILD": "OWN"}).fillna("OTHER")
    own = ws[(ws.rel == "OWN") & (ws.lives_in_home == "Y")].copy()
    own["birth"] = own.child_birth_date.map(parse_birth)
    own = own.drop_duplicates(["no", "child_initials", "birth"])
    births = own.groupby("no")["birth"].apply(lambda s: sorted(s)).to_dict()

    # home-months, January 2016 to December 2026
    m0, m1 = 2016 * 12, 2026 * 12 + 11
    recs = []
    stock = []
    ksum_l, kcl_l = [0], [0]
    for r in lic.itertuples():
        first = dt.date.fromisoformat(r.first_licensed_date)
        close = dt.date.fromisoformat(r.closure_effective_date) if r.closure_effective_date else None
        by = int(r.provider_birth_year) if r.provider_birth_year else None
        bs = births.get(r.licence_no, [])
        lm = mindex(first)
        cm = mindex(close) if close else None
        rng_months = list(range(max(lm + 1, m0), (min(cm, m1) if cm is not None else m1) + 1))
        if close is None:
            stock.append((r.licence_no, lm, by, bs))
        june_years = set()
        for mm in rng_months:
            y, mo = divmod(mm, 12)
            mo += 1
            mend = dt.date(y + (mo == 12), mo % 12 + 1, 1) - dt.timedelta(days=1)
            yb = [b for b in bs if b <= mend]
            k = 1 if (yb and kyear(max(yb)) == y and mo in (6, 7, 8)) else 0
            if k and mo == 6:
                june_years.add(y)
                ksum_l[0] += 1
            if k and cm == mm and y in june_years:
                kcl_l[0] += 1
            recs.append((min((mm - lm) // 12, 10), int(by is not None and y - by >= 63), int(by is not None and y - by >= 68),
                         int(mo in (6, 7, 8)), k, int(cm == mm)))
    H = pd.DataFrame(recs, columns=["tb", "a63", "a68", "su", "k", "ev"])
    ksum, kcl = ksum_l[0], kcl_l[0]
    say("V home-months", len(H) > 150000, "%d rows, %d closures" % (len(H), H.ev.sum()))
    cellsdf = H.groupby(["tb", "a63", "a68", "su", "k"]).agg(ev=("ev", "sum"), n=("ev", "size")).reset_index()

    def X_of(df):
        cols = [np.ones(len(df))] + [(df.tb.values == b).astype(float) for b in range(1, 11)]
        cols += [df.a63.values.astype(float), df.a68.values.astype(float), df.su.values.astype(float), df.k.values.astype(float)]
        return np.column_stack(cols)

    X = X_of(cellsdf)
    keep = X.std(axis=0) > 0
    keep[0] = True
    beta = irls_poisson(X[:, keep], cellsdf.ev.values.astype(float), cellsdf.n.values.astype(float))
    tot = 0.0
    nk = 0
    for no, lm, by, bs in stock:
        s = 1.0
        flag = False
        for mo in range(1, 13):
            mm = 2027 * 12 + mo - 1
            mend = dt.date(2027 + (mo == 12), mo % 12 + 1, 1) - dt.timedelta(days=1)
            yb = [b for b in bs if b <= mend]
            k = 1 if (yb and kyear(max(yb)) == 2027 and mo in (6, 7, 8)) else 0
            flag = flag or bool(k)
            row = pd.DataFrame([((min((mm - lm) // 12, 10)), int(by is not None and 2027 - by >= 63),
                                 int(by is not None and 2027 - by >= 68), int(mo in (6, 7, 8)), k)],
                               columns=["tb", "a63", "a68", "su", "k"]) if False else None
            x = np.array([1.0] + [float(min((mm - lm) // 12, 10) == b) for b in range(1, 11)] +
                         [float(by is not None and 2027 - by >= 63), float(by is not None and 2027 - by >= 68),
                          float(mo in (6, 7, 8)), float(k)])[keep]
            s *= 1 - min(np.exp(x @ beta), 0.99)
        tot += 1 - s
        nk += int(flag)
    call = int(round(tot / 50) * 50)
    say("V the call to the nearest 50", call == figs["call"]["closures_2027"], "%.2f -> %d" % (tot, call))
    say("V homes licensed on 31 December 2026", len(stock) == figs["cc1"], "%d" % len(stock))
    say("V homes in their youngest's kindergarten summer in 2027", nk == figs["cc2"], "%d" % nk)
    # the record's kindergarten summers: homes in a flagged June, and those closing in flagged months
    say("V record's kindergarten summers and their closures", (ksum, kcl) == (figs["cc3"], figs["cc5"]), "%d of %d" % (kcl, ksum))
    # the tenure-only route
    tl = H.groupby("tb").ev.mean().to_dict()
    tt = 0.0
    for no, lm, by, bs in stock:
        s = 1.0
        for mo in range(1, 13):
            mm = 2027 * 12 + mo - 1
            s *= 1 - tl[min((mm - lm) // 12, 10)]
        tt += 1 - s
    say("V the tenure-only route", int(round(tt / 50) * 50) == figs["alt_tenure"], "%.1f" % tt)

    # ---- asks ------------------------------------------------------------------------------------------------------
    cty = pd.read_csv(os.path.join(target, "county_region_lookup.csv"), dtype=str)
    reg_f = dict(zip(cty.county_fips, cty.region_no.astype(int)))
    reg_n = dict(zip(cty.county_name.str.replace(" County", "").str.upper(), cty.region_no.astype(int)))
    names = dict(zip(cty.region_no.astype(int), cty.region_name))
    slug = {r: "R%d_%s" % (r, names[r].split()[-1]) for r in names}
    cc = pd.read_csv(os.path.join(target, "ccap_payment_lines_2026.csv"), dtype=str)
    cc = cc[cc.provider_type.isin(["SFCC", "LFCC"]) & (cc.remittance_status != "VOID")]
    days = []
    for r in cc.itertuples():
        a = max(dt.date.fromisoformat(r.service_from), dt.date(2026, 1, 1))
        b = min(dt.date.fromisoformat(r.service_to), dt.date(2026, 12, 31))
        for d in pd.bdate_range(a, b) if a <= b else []:
            days.append((r.child_no, d.date(), reg_f[r.county_fips]))
    D = pd.DataFrame(days, columns=["child", "day", "reg"]).drop_duplicates(["child", "day"])
    D["half"] = np.where(pd.to_datetime(D.day).dt.month <= 6, "H1", "H2")
    A = D.groupby("half").size().to_dict()
    say("V ask A statewide", figs["a_statewide"] == {"H1": A["H1"], "H2": A["H2"]}, str(A))
    rm = cc[cc.pay_run_date.str.startswith("2026")].drop_duplicates(["remittance_no", "pay_run_date", "county_fips"]).copy()
    rm["amt"] = rm.remittance_amount.astype(float)
    rm["reg"] = rm.county_fips.map(reg_f)
    rm["q"] = (rm.pay_run_date.str[5:7].astype(int) - 1) // 3 + 1
    B = rm.groupby(["reg", "q"]).amt.sum().to_dict()
    for r in names:
        say("V ask B %s" % slug[r], figs["b_" + slug[r]] == {"Q%d" % q: int(round(B[(r, q)])) for q in (1, 2, 3, 4)})

    # ---- every format read back -----------------------------------------------------------------------------------
    import openpyxl
    import pypdf
    wb = openpyxl.load_workbook(os.path.join(target, "startup_grant_awards_2023.xlsx"), read_only=True)
    say("V grant workbook reads", wb["Awards"].max_row > 1000, "")
    wb2 = openpyxl.load_workbook(os.path.join(target, "market_rate_survey_2025.xlsx"), read_only=True)
    say("V market rate workbook reads", wb2.sheetnames == ["Weekly full-time rates"], "")

    def ptext(f):
        return " ".join(p.extract_text() or "" for p in pypdf.PdfReader(os.path.join(target, f)).pages).replace("\n", " ")
    pol = ptext("ccap_provider_payment_policy_2026.pdf")
    say("V policy: the school-age rule reads back", "turns five on or before 1 September" in pol, "")
    say("V policy: positive control", "Providers are paid monthly in arrears" in pol, "")
    rules = ptext("family_child_care_licensing_rules_2025.pdf")
    say("V rules: two licence types", "Family child care homes are licensed in two types" in rules, "")
    say("V rules: worksheet children", "capacity worksheet" in rules, "")
    sup = ptext("child_care_supply_report_2026q3.pdf")
    say("V supply report reads", "start-up grant round" in sup.lower() or "Start-Up Grant" in sup, "")
    txt = open(os.path.join(target, "data_dictionary.txt"), encoding="utf-8").read()
    say("V dictionary reads", "capacity_worksheet_children_2016_2026.csv" in txt, "")
    gold = os.path.join(task, "golden", "fcc_closures_2027.xlsx")
    if os.path.exists(gold):
        g = openpyxl.load_workbook(gold, read_only=True)
        sh = g["Forecast 2027"]
        vals = [c for row in sh.iter_rows(values_only=True) for c in row if isinstance(c, (int, float))]
        say("V golden workbook carries the call", figs["call"]["closures_2027"] in vals and figs["cc1"] in vals, "")
    print("verify: %d failed" % len(BAD))
    return 1 if BAD else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(HERE), "target")))
