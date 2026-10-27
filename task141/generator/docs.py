"""The pack's documents: SB 26-118 as enrolled, the State Jail Standards Board's instructions for the 2027 plans, the
jail, court and office code lists, and the Department's annual report for fiscal 2025 (its tables computed from the
shipped registers and petitions)."""
from __future__ import annotations

import datetime as dt

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import LETTER                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import inch                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import common as K    # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/"
for _n, _f in (("Carlito", "crosextra/Carlito-Regular.ttf"), ("Carlito-Bold", "crosextra/Carlito-Bold.ttf"),
               ("Carlito-Italic", "crosextra/Carlito-Italic.ttf"),
               ("LiberationSerif", "liberation/LiberationSerif-Regular.ttf"),
               ("LiberationSerif-Bold", "liberation/LiberationSerif-Bold.ttf"),
               ("LiberationSerif-Italic", "liberation/LiberationSerif-Italic.ttf")):
    pdfmetrics.registerFont(TTFont(_n, FONT_DIR + _f))
SANS, SANS_B, SANS_I = "Carlito", "Carlito-Bold", "Carlito-Italic"
SERIF, SERIF_B, SERIF_I = "LiberationSerif", "LiberationSerif-Bold", "LiberationSerif-Italic"
INK = colors.HexColor("#1b1f27")
GREY = colors.HexColor("#555b66")
RULE = colors.HexColor("#9aa3ad")
PALE = colors.HexColor("#eceff3")

BODY = ParagraphStyle("body", fontName=SERIF, fontSize=10.8, leading=14.4, textColor=INK, spaceAfter=6)
CLAUSE = ParagraphStyle("clause", parent=BODY, leftIndent=26, firstLineIndent=-26)
H1 = ParagraphStyle("h1", fontName=SERIF_B, fontSize=13.5, leading=17, textColor=INK, spaceAfter=3, alignment=1)
H1S = ParagraphStyle("h1s", fontName=SERIF, fontSize=10.4, leading=13.2, textColor=INK, spaceAfter=10, alignment=1)
H2 = ParagraphStyle("h2", fontName=SERIF_B, fontSize=11, leading=14, textColor=INK, spaceBefore=8, spaceAfter=3)
NOTE = ParagraphStyle("note", fontName=SERIF_I, fontSize=9.2, leading=11.8, textColor=GREY, spaceBefore=6)
SB = ParagraphStyle("sb", fontName=SANS, fontSize=10.4, leading=13.6, textColor=INK, spaceAfter=6)
SBC = ParagraphStyle("sbc", parent=SB, leftIndent=24, firstLineIndent=-24)
SH1 = ParagraphStyle("sh1", fontName=SANS_B, fontSize=15.5, leading=19, textColor=colors.HexColor("#1f3a5a"), spaceAfter=2)
SH1S = ParagraphStyle("sh1s", fontName=SANS, fontSize=9.8, leading=12.6, textColor=GREY, spaceAfter=10)
SH2 = ParagraphStyle("sh2", fontName=SANS_B, fontSize=11.4, leading=14.2, textColor=INK, spaceBefore=9, spaceAfter=3)
CELL = ParagraphStyle("cell", fontName=SANS, fontSize=9, leading=11, textColor=INK)
CELLR = ParagraphStyle("cellr", parent=CELL, alignment=2)
CELLB = ParagraphStyle("cellb", fontName=SANS_B, fontSize=9, leading=11, textColor=INK)
CELLBR = ParagraphStyle("cellbr", parent=CELLB, alignment=2)


def n0(x):
    return "{:,}".format(int(round(x)))


def _footer(left, right):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont(SANS, 7.8)
        canvas.setFillColor(GREY)
        canvas.drawString(0.9 * inch, 0.55 * inch, left)
        canvas.drawRightString(LETTER[0] - 0.9 * inch, 0.55 * inch, "%s  |  page %d" % (right, doc.page))
        canvas.restoreState()
    return draw


def _doc(path, title, author):
    return SimpleDocTemplate(path, pagesize=LETTER, leftMargin=1.0 * inch, rightMargin=1.0 * inch, topMargin=0.85 * inch,
                             bottomMargin=0.9 * inch, title=title, author=author, subject="", creator="", producer="")


def _table(rows, widths, head=True, right_from=1, total=False):
    data = []
    for i, r in enumerate(rows):
        hd = head and i == 0
        tl = total and i == len(rows) - 1
        out = []
        for j, c in enumerate(r):
            num = j >= right_from and i > 0
            st = (CELLBR if num else CELLB) if (hd or tl) else (CELLR if num else CELL)
            if hd and j >= right_from:
                st = CELLBR
            out.append(Paragraph(str(c), st))
        data.append(out)
    t = Table(data, colWidths=widths, hAlign="LEFT")
    st = [("LINEABOVE", (0, 0), (-1, 0), 0.6, RULE), ("LINEBELOW", (0, 0), (-1, 0), 0.6, RULE),
          ("LINEBELOW", (0, -1), (-1, -1), 0.6, RULE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
          ("TOPPADDING", (0, 0), (-1, -1), 2.4), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.4)]
    if head:
        st.append(("BACKGROUND", (0, 0), (-1, 0), PALE))
    if total:
        st.append(("LINEABOVE", (0, -1), (-1, -1), 0.6, RULE))
    t.setStyle(TableStyle(st))
    return t


# ---------------------------------------------------------------------------------------------------- SB 26-118
def write_act(path):
    s = [Paragraph("SENATE BILL 26-118", H1),
         Paragraph("ENROLLED", H1S),
         Paragraph("AN ACT concerning the length of felony probation terms.", H1S),
         Paragraph("<i>Be it enacted by the General Assembly of the State:</i>", BODY)]

    def sec(n, title, text):
        s.append(Paragraph("<b>SECTION %d. %s.</b> %s" % (n, title, text), BODY))

    def sub(text):
        s.append(Paragraph(text, CLAUSE))
    sec(1, "Short title and legislative declaration", "")
    sub("(1)&nbsp;&nbsp;This Act may be cited as the Felony Probation Term Act.")
    sub("(2)&nbsp;&nbsp;The General Assembly finds and declares that a felony probation term should end within a period "
        "the probationer and the court can foresee at sentencing, that the State's felony probation terms run longer "
        "than those of its neighboring states, and that a limit of 24 months on supervision serves the purpose of this "
        "Act.")
    sec(2, "Definitions", "In this Act:")
    sub("(1)&nbsp;&nbsp;\"Felony probation term\" means a term of probation imposed on conviction of a felony, with the "
        "conditions the court orders, including any period of confinement in a county jail the court orders as a "
        "condition of the term.")
    sub("(2)&nbsp;&nbsp;\"Supervision\" on a term means supervision by a county probation department. Supervision on a "
        "term begins on the day the probationer is placed under the department's supervision on that term, which for a "
        "term with a confinement condition is the day of release from the confinement.")
    sub("(3)&nbsp;&nbsp;\"Scheduled expiration\" means the day on which a term is to expire under the order imposing it "
        "as amended by any later order of the court.")
    sec(3, "Terms imposed on or after the effective date", "A felony probation term imposed on or after the effective "
        "date of this Act may not provide for supervision for more than 24 months, and expires no later than the day 24 "
        "months after supervision on it begins.")
    sec(4, "Terms in effect", "A felony probation term in effect on the effective date of this Act expires on the "
        "earlier of (a) its scheduled expiration and (b) the later of the day 24 months after supervision on it began and "
        "31 March 2027.")
    sec(5, "Extension and tolling", "A court may extend or toll a felony probation term as before, but no order made "
        "on or after the effective date of this Act may set the term's expiration later than the day sections 3 and 4 "
        "allow.")
    sec(6, "Pending proceedings", "Proceedings pending on the day a term ends under section 4 continue.")
    sec(7, "Misdemeanor terms", "This Act does not apply to a term of probation imposed on conviction of a misdemeanor.")
    sec(8, "Effective date", "This Act takes effect on 1 January 2027.")
    sec(9, "Safety clause", "The General Assembly finds, determines and declares that this Act is necessary for the "
        "preservation of the public peace, health and safety.")
    s.append(Spacer(1, 10))
    s.append(Paragraph("Passed the Senate 21 April 2026; passed the House 6 May 2026. Approved by the Governor 14 May "
                       "2026.", NOTE))
    doc = _doc(path, "SB 26-118, Felony Probation Term Act (enrolled)", "Legislative Services")
    f = _footer("SB 26-118, enrolled", "Felony Probation Term Act")
    doc.build(s, onFirstPage=f, onLaterPages=f)


# ---------------------------------------------------------------------------------------------------- the Board
STANDARDS = (("Vessey County", 48, 160, "12 May 2020"), ("Brannock County", 56, 190, "2 February 2022"),
             ("Okelan County", K.POD_BEDS, K.POD_SIZE, "3 March 2021"), ("Ostby County", 64, 200, "9 June 2021"),
             ("Quillan County", 40, 150, "17 November 2023"))


def write_plan(path):
    s = [Paragraph("State Jail Standards Board", SH1S), Paragraph("2027 Jail Operations Plans: instructions to counties",
                                                                  SH1),
         Paragraph("Issued 8 July 2026 under the Board's rule 12 (county jail operations plans). Replaces the "
                   "instructions for the 2026 plans.", SH1S)]

    def head(t):
        s.append(Paragraph(t, SH2))

    def cl(n, t):
        s.append(Paragraph("<b>%s</b>&nbsp;&nbsp;%s" % (n, t), SBC))
    head("1. Filing")
    cl("1.1", "Each county files a Jail Operations Plan for calendar 2027 (the plan year), adopted by its board of "
              "supervisors, with the Board by 1 December 2026.")
    cl("1.2", "The plan is built on the county's records to the end of 30 September 2026 (the extract).")
    cl("1.3", "Dates and times in a plan, and county time in these instructions, are Central time.")
    head("2. Violator housing")
    cl("2.1", "Under the State's 2019 realignment of revocation terms a person committed on the revocation of a felony "
              "probation term serves the commitment in the jail of the county that supervised the term.")
    cl("2.2", "The plan states the felony probation revocation admissions the county projects for the plan year and "
              "the violator housing it will staff for them under the housing standard the county has on file with the "
              "Board (Annex). The Board approves the plan's violator housing staffing on that projection.")
    head("3. Admissions")
    cl("3.1", "An admission is a person's commitment to the county jail on the revocation of a probation term the "
              "county probation department supervises. It is counted once, on the day the court orders the "
              "commitment, however many of the person's terms the court revoked that day.")
    cl("3.2", "An admission is a felony probation revocation admission when a term revoked is a felony probation term. "
              "Holds on a violation warrant before the court decides, and holds for another state, are not admissions.")
    head("4. The projection")
    cl("4.1", "Project the plan year's felony probation revocation admissions under the law in force in the plan year, "
              "with the courts, the probation department and the sheriff acting as the county's records to the "
              "extract show them acting.")
    cl("4.2", "The projection counts every admission the plan year will see, including admissions on terms whose "
              "supervision begins after the extract.")
    cl("4.3", "State the violator housing units the projection calls for under the county's standard (Annex).")
    head("5. Violator housing in use")
    cl("5.1", "Report the violator bed-days of each of the four quarters to the extract, felony and misdemeanor.")
    cl("5.2", "A violator bed-day is counted for each person in the jail at 00:00 county time serving a commitment "
              "the court has ordered on the revocation of a probation term, in any of the jail's buildings, and is "
              "dated to the day that begins at that midnight. A person serving a commitment on weekends is counted at "
              "each midnight spent in the jail.")
    cl("5.3", "A bed-day is a felony bed-day when the commitment revokes a felony probation term; a person held at "
              "that midnight on commitments revoking a felony term and a misdemeanor term is counted once, as a "
              "felony bed-day. Every other violator bed-day is a misdemeanor bed-day. Holds on a violation warrant "
              "before the court decides, and holds for another state, are not violator bed-days.")
    head("6. Felony probationers on supervision")
    cl("6.1", "Report, for each office supervising them at the extract, the felony probationers the county probation "
              "department supervised at the extract, by time on supervision at the extract in completed years from the "
              "day supervision began: under one year, one year to under two years, two years to under three years, "
              "and three years or more.")
    head("7. Changes in law")
    cl("7.1", "Senate Bill 26-118, the Felony Probation Term Act, takes effect on 1 January 2027.")
    s.append(Paragraph("Annex. Violator housing standards on file with the Board, 1 July 2026", SH2))
    rows = [["County", "Violator housing unit", "Standard", "Filed"]]
    for cty, beds, adm, filed in STANDARDS:
        rows.append([cty, "pod of %d beds" % beds, "one unit for each %d felony probation revocation admissions "
                     "projected for the plan year, or part of %d" % (adm, adm), filed])
    s.append(_table(rows, [1.15 * inch, 1.25 * inch, 3.0 * inch, 1.1 * inch], right_from=9))
    s.append(Paragraph("A county changing its standard files the new standard with its plan.", NOTE))
    doc = _doc(path, "2027 Jail Operations Plans: instructions to counties", K.BOARD)
    f = _footer("SJSB instructions for the 2027 Jail Operations Plans", "State Jail Standards Board")
    doc.build(s, onFirstPage=f, onLaterPages=f)


# ---------------------------------------------------------------------------------------------------- annual report
RISK_LEVELS = (("LOW", "Low (CSRS 0 to 7)", 0, 7), ("MOD", "Moderate (CSRS 8 to 13)", 8, 13),
               ("HIGH", "High (CSRS 14 to 20)", 14, 20))


def report_figures(cms1, cms2, pets):
    """the Department's fiscal 2025 tables (1 July 2024 to 30 June 2025), computed from the shipped rows"""
    a, b = dt.date(2024, 7, 1), dt.date(2025, 6, 30)
    leg = {o[2]: o[0] for o in K.OFFICES}

    def mdy(s):
        return dt.datetime.strptime(s, "%m/%d/%Y").date() if s else None
    terms = []
    for r in cms1:
        terms.append(dict(lvl=r["LEVEL"], risk=int(r["RISK_SCORE"]) // 2, sup=mdy(r["SUP_BEGIN"]),
                          closed=mdy(r["CLOSE_DATE"]),
                          reason={"EXP": "EXPIRED", "REV": "REVOKED", "ETERM": "EARLY TERMINATION"}[r["CLOSE_CODE"]]))
    for r in cms2:
        terms.append(dict(lvl=r["offense_level"][0], risk=int(r["csrs_score"]),
                          sup=dt.date.fromisoformat(r["supervision_begin"]),
                          closed=dt.date.fromisoformat(r["closed_date"]) if r["closed_date"] else None,
                          reason=r["closure_reason"]))
    cl = {k: {"F": 0, "M": 0} for k, _, _, _ in RISK_LEVELS}
    began = {"F": 0, "M": 0}
    closed = {(lv, rs): 0 for lv in "FM" for rs in ("EXPIRED", "REVOKED", "EARLY TERMINATION")}
    for t in terms:
        if t["sup"] <= b and (t["closed"] is None or t["closed"] > b):
            cl[[k for k, _, lo, hi in RISK_LEVELS if lo <= t["risk"] <= hi][0]][t["lvl"]] += 1
        if a <= t["sup"] <= b:
            began[t["lvl"]] += 1
        if t["closed"] is not None and a <= t["closed"] <= b:
            closed[(t["lvl"], t["reason"])] += 1
    filed = {"F": 0, "M": 0}
    disp = {(lv, d): 0 for lv in "FM" for d in ("REVOKED", "CONTINUED", "CONTINUED AND EXTENDED", "DISMISSED")}
    for p in pets:
        lv = "F" if "-CF-" in p["case_number"] else "M"
        if a <= dt.date.fromisoformat(p["filed_date"]) <= b:
            filed[lv] += 1
        if p["disposition_date"] and a <= dt.date.fromisoformat(p["disposition_date"]) <= b:
            disp[(lv, p["disposition"])] += 1
    return dict(caseload=cl, began=began, closed=closed, filed=filed, disp=disp)


def write_report(path, F):
    s = [Paragraph("Okelan County Adult Probation Department", SH1S), Paragraph("Annual report, fiscal year 2025", SH1),
         Paragraph("1 July 2024 to 30 June 2025. Prepared by the Research and Data Unit for the Chief Probation "
                   "Officer's report to the County Board of Supervisors. Counts are from CMS 2 and the Clerk of the "
                   "Circuit Court's case management system as at 31 August 2025.", SH1S),
         Paragraph("The year", SH2)]
    cl = F["caseload"]
    totf = sum(v["F"] for v in cl.values())
    totm = sum(v["M"] for v in cl.values())
    s.append(Paragraph("At the end of the fiscal year the Department supervised %s felony and %s misdemeanor "
                       "probationers. Officers filed %s violation petitions on felony terms and %s on misdemeanor terms, "
                       "and the Circuit Court revoked %s felony and %s misdemeanor terms."
                       % (n0(totf), n0(totm), n0(F["filed"]["F"]), n0(F["filed"]["M"]), n0(F["disp"][("F", "REVOKED")]),
                          n0(F["disp"][("M", "REVOKED")])), SB))
    s.append(Paragraph("Table 1. Probationers on supervision at 30 June 2025, by risk level at intake", SH2))
    rows = [["Risk level", "Felony", "Misdemeanor", "Total"]]
    for code, name, _, _ in RISK_LEVELS:
        rows.append([name, n0(cl[code]["F"]), n0(cl[code]["M"]), n0(cl[code]["F"] + cl[code]["M"])])
    rows.append(["All levels", n0(totf), n0(totm), n0(totf + totm)])
    s.append(_table(rows, [2.0 * inch, 1.1 * inch, 1.1 * inch, 1.1 * inch], total=True))
    s.append(Paragraph("Table 2. Violation petitions disposed in the year, by disposition", SH2))
    rows = [["Disposition", "Felony terms", "Misdemeanor terms"]]
    for d, lab in (("REVOKED", "Revoked"), ("CONTINUED", "Continued"), ("CONTINUED AND EXTENDED", "Continued and extended"),
                   ("DISMISSED", "Dismissed")):
        rows.append([lab, n0(F["disp"][("F", d)]), n0(F["disp"][("M", d)])])
    rows.append(["All dispositions", n0(sum(F["disp"][("F", d)] for d in ("REVOKED", "CONTINUED", "CONTINUED AND EXTENDED",
                                                                          "DISMISSED"))),
                 n0(sum(F["disp"][("M", d)] for d in ("REVOKED", "CONTINUED", "CONTINUED AND EXTENDED", "DISMISSED")))])
    s.append(_table(rows, [2.4 * inch, 1.3 * inch, 1.5 * inch], total=True))
    s.append(Paragraph("Table 3. Terms closed in the year, by reason", SH2))
    rows = [["Reason", "Felony", "Misdemeanor"]]
    for rs, lab in (("EXPIRED", "Expired"), ("REVOKED", "Revoked"), ("EARLY TERMINATION", "Early termination")):
        rows.append([lab, n0(F["closed"][("F", rs)]), n0(F["closed"][("M", rs)])])
    rows.append(["All closures", n0(sum(F["closed"][("F", rs)] for rs in ("EXPIRED", "REVOKED", "EARLY TERMINATION"))),
                 n0(sum(F["closed"][("M", rs)] for rs in ("EXPIRED", "REVOKED", "EARLY TERMINATION")))])
    s.append(_table(rows, [2.4 * inch, 1.3 * inch, 1.5 * inch], total=True))
    s.append(Paragraph("Notes", SH2))
    s.append(Paragraph("A probationer on supervision is counted once for each open term. A petition is counted in the "
                       "year it was filed and its disposition in the year the court entered it, so the two tables cover "
                       "different petitions. Early termination is granted on the officer's motion after eighteen months "
                       "of compliance on a felony term and six on a misdemeanor term.", SB))
    s.append(Paragraph("Risk levels are the Community Supervision Risk Screen score taken at intake; terms opened in "
                       "CMS 1 are shown on the screen's current scale.", SB))
    doc = _doc(path, "Adult Probation Department annual report, fiscal year 2025", K.PEOPLE["research"])
    f = _footer("APD annual report FY2025", "Okelan County Adult Probation Department")
    doc.build(s, onFirstPage=f, onLaterPages=f)


# ---------------------------------------------------------------------------------------------------- code lists
JMS_CODES = [
    ("classification", "PROB-COND", "Probation condition confinement"),
    ("classification", "PV-HOLD", "Probation violation warrant hold"),
    ("classification", "PV-FEL", "Probation violation commitment, felony probationer"),
    ("classification", "PV-MIS", "Probation violation commitment, misdemeanor probationer"),
    ("classification", "PV-OOS", "Interstate Compact hold for another state"),
    ("facility", "MAIN", "Main building"),
    ("facility", "ANNEX", "Harlan Road annex"),
    ("release_reason", "SENTENCE COMPLETE", "Sentence complete"),
    ("release_reason", "RETURNED TO SUPERVISION", "Returned to probation supervision"),
    ("release_reason", "RECLASSIFIED", "Reclassified"),
    ("release_reason", "TRANSFER TO ANNEX", "Transferred to the annex"),
    ("release_reason", "SCHEDULED RELEASE", "Weekend sentence period ended"),
    ("release_reason", "RELEASED ON COURT ORDER", "Released on the court's order"),
    ("release_reason", "RELEASED TO SENDING STATE", "Released to the sending state"),
    ("release_reason", "CONVERTED TO NEW JMS", "Converted to the new jail management system"),
]
OFFICE_ROWS = [("CEN", "C", "Central", "3 January 1994"), ("NGT", "N", "Northgate", "2 July 2001"),
               ("RVS", "R", "Riverside", "2 July 2001"), ("EFD", "", "Eastfield", "2 June 2025")]
COURT_SHEETS = (
    ("Allegations", ["Code", "Allegation"], [
        ("LAW", "New law violation: arrest on or charge with a new offense"),
        ("RPT", "Failure to report to the supervising officer"),
        ("DRG", "Positive, adulterated or missed drug or alcohol test"),
        ("TRT", "Failure to attend or complete ordered treatment"),
        ("PAY", "Failure to pay ordered restitution, fines or fees"),
        ("CUR", "Curfew or electronic monitoring violation"),
        ("CNT", "Contact with a person the court ordered no contact with"),
        ("ABS", "Absconding: whereabouts unknown to the supervising officer")]),
    ("Dispositions", ["Disposition", "Meaning"], [
        ("REVOKED", "Probation revoked, commitment ordered"),
        ("CONTINUED", "Probation continued"),
        ("CONTINUED AND EXTENDED", "Probation continued, term extended"),
        ("DISMISSED", "Petition dismissed")]),
    ("Process", ["Process", "Meaning"], [
        ("WARRANT", "Violation warrant issued"),
        ("SUMMONS", "Summons issued")]),
    ("Term orders", ["Order type", "Meaning"], [
        ("TOLLING", "Term tolled while the probationer absconded"),
        ("EXTENSION", "Term extended"),
        ("EARLY TERMINATION", "Term terminated early")]),
    ("Served as", ["Value", "Meaning"], [
        ("CONSECUTIVE", "Served consecutively"),
        ("WEEKENDS", "Served on weekends")]),
    ("Case numbers", ["Type", "Meaning"], [
        ("CF", "Felony"),
        ("CM", "Misdemeanor")]),
    ("Offense classes", ["Class", "Meaning", "Longest probation term (months)"], [
        ("F2", "Class 2 felony", 96),
        ("F3", "Class 3 felony", 84),
        ("F4", "Class 4 felony", 60),
        ("F5", "Class 5 felony", 48),
        ("F6", "Class 6 felony", 36),
        ("DF4", "Drug felony, class 4", 24),
        ("M1", "Class 1 misdemeanor", 24),
        ("M2", "Class 2 misdemeanor", 12)]),
)
CLASS_NOTE = ("Longest probation term: the longest term of probation the Criminal Code allows for the class. No order "
              "imposing or extending a term may set it longer.")


def write_jms_codes(path):
    K.write_csv(path, ["field", "code", "description"], [list(x) for x in JMS_CODES], crlf=True)


def write_office_codes(path):
    K.write_csv(path, ["office_code", "cms1_code", "office_name", "in_use_from"], [list(x) for x in OFFICE_ROWS],
                crlf=True)


def write_court_codes(path):
    wb = openpyxl.Workbook()
    for i, (title, cols, rows) in enumerate(COURT_SHEETS):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append(["Clerk of the Circuit Court, Okelan County: code table, %s" % title.lower()])
        ws["A1"].font = Font(name="Calibri", bold=True, size=11)
        ws.append([])
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=3, column=c)
            cell.font = Font(name="Calibri", bold=True)
            cell.fill = PatternFill("solid", fgColor="DDE4EC")
            cell.alignment = Alignment(vertical="top")
        for r in rows:
            ws.append(list(r))
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 84 if len(cols) == 2 else 40
        if len(cols) == 3:
            ws.column_dimensions["C"].width = 32
            ws.append([])
            ws.append([CLASS_NOTE])
            ws.cell(ws.max_row, 1).font = Font(name="Calibri", italic=True, size=9)
        ws.freeze_panes = "A4"
    wb.save(path)
