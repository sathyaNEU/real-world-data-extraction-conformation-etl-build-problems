"""task122 generator: the filed documents. The phase 2b protocol as amended, its statistical analysis plan, the draft
phase 3 synopsis, the central laboratory's manual, the CRO's data dictionary, the site list and the export manifest."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess

ORG = P.ORG


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["BodyText"], fontName="Helvetica", fontSize=9.4, leading=12.8, spaceAfter=4)
    h = ParagraphStyle("h", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, spaceBefore=8, spaceAfter=3)
    t = ParagraphStyle("t", parent=ss["Title"], fontName="Helvetica-Bold", fontSize=14, alignment=0, spaceAfter=4)
    small = ParagraphStyle("s", parent=body, fontSize=8, leading=10.5, textColor=colors.HexColor("#444444"))
    return body, h, t, small


def _footer(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(20 * mm, 10 * mm, text)
        canvas.drawRightString(190 * mm, 10 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _pdf(path, title, author, footer, blocks, stamp):
    body, h, t, small = _styles()
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                            bottomMargin=18 * mm, title=title, author=author, creator=author, invariant=1)
    flow = []
    for kind, val in blocks:
        if kind == "title":
            flow.append(Paragraph(val, t))
        elif kind == "h":
            flow.append(Paragraph(val, h))
        elif kind == "p":
            flow.append(Paragraph(val, body))
        elif kind == "small":
            flow.append(Paragraph(val, small))
        elif kind == "sp":
            flow.append(Spacer(1, val))
        elif kind == "page":
            flow.append(PageBreak())
        elif kind == "table":
            rows, widths = val
            cell = ParagraphStyle("c", parent=body, fontName="Helvetica", fontSize=7.8, leading=9.6, spaceAfter=0)
            cellb = ParagraphStyle("cb", parent=cell, fontName="Helvetica-Bold")
            rows = [[Paragraph(str(x), cellb if r == 0 else cell) for x in row] for r, row in enumerate(rows)]
            tb = Table(rows, colWidths=widths, repeatRows=1)
            tb.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#9a9a9a")),
                                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6eaef")),
                                    ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            flow.append(tb)
    doc.build(flow, onFirstPage=_footer(footer), onLaterPages=_footer(footer))
    postprocess.scrub(path, ORG, stamp)


def _visit_table():
    rows = [["Visit", "Visit number", "Week", "Study day", "Window (days)"]]
    for v, n, day, win in P.VISITS:
        wk = {"SCREENING": "-2", "DAY 1": "0"}[v] if v in ("SCREENING", "DAY 1") else v.split()[1]
        rows.append([v.title().replace("Follow-Up", "follow-up"), str(n), wk, str(day), "+/- %d" % win if win else "none"])
    return rows


# ------------------------------------------------------------------------------------------------ protocol
def protocol(path):
    ep = [["Objective", "End point"],
          ["Primary: the effect of VMK-4410 on albuminuria", "Change in urine albumin-to-creatinine ratio (UACR) from baseline to week 24 (ratio to baseline)"],
          ["Secondary: the effect on kidney function on treatment", "Change in eGFR from baseline to week 52"],
          ["Secondary: the effect on kidney function after treatment stops", "Change in eGFR from baseline to the week-56 follow-up visit, four weeks after the last dose"],
          ["Safety", "Adverse events; serum potassium; serum sodium; blood pressure"],
          ["Exploratory", "Trough plasma concentrations of VMK-4410 at weeks 4, 12, 24 and 52"]]
    blocks = [
        ("title", "Clinical study protocol VMK-4410-202"),
        ("p", "<b>A 52-week, randomised, double-blind, placebo-controlled, dose-ranging phase 2b study of VMK-4410 in adults "
              "with chronic kidney disease and albuminuria</b>"),
        ("small", "Sponsor: Varnmark Therapeutics A/S. Protocol version 3.0 (amendment 2), 10 June 2024. Investigational "
                  "product: VMK-4410, an oral aldosterone synthase inhibitor, tablets of 5, 10, 20 and 40 mg and matching "
                  "placebo. This document replaces version 2.0 of 5 February 2024."),
        ("h", "Amendment history"),
        ("table", ([["Version", "Date", "Main changes"],
                    ["1.0", "20 November 2023", "Original protocol"],
                    ["2.0 (amendment 1)", "5 February 2024", "Sites added in Spain and the United Kingdom; the Madrid hub of the central laboratory added for them"],
                    ["3.0 (amendment 2)", "10 June 2024", "Potassium rules restated in section 7 (permanent stop above 6.5 mmol/L or above 6.0 mmol/L a second time); repeat potassium tests at a local laboratory permitted; urine collection added at week 4"]],
                   [30 * mm, 32 * mm, 108 * mm])),
        ("h", "1. Background"),
        ("p", "Aldosterone contributes to kidney injury in chronic kidney disease through inflammation and fibrosis, and "
              "its effects are not removed by ACE inhibitors or angiotensin receptor blockers. VMK-4410 inhibits aldosterone "
              "synthase. In phase 1 it lowered plasma aldosterone in a dose-dependent way and raised serum potassium at the "
              "higher doses. Drugs acting on the renin-angiotensin-aldosterone system lower eGFR in the first weeks of "
              "treatment, a fall that reverses when treatment stops."),
        ("h", "2. Objectives and end points"),
        ("table", (ep, [70 * mm, 100 * mm])),
        ("h", "3. Design"),
        ("p", "Patients are randomised 1:1:1:1:1 to placebo or VMK-4410 5 mg, 10 mg, 20 mg or 40 mg once daily for 52 weeks, "
              "through the interactive response system, in blocks within four strata: screening eGFR below 45 or 45 and above "
              "(mL/min/1.73 m2) and screening UACR below 1,000 or 1,000 mg/g and above. About 750 patients are randomised at "
              "about 50 sites in Denmark, Sweden, Germany, Poland, Spain and the United Kingdom. Patients, investigators and "
              "the sponsor's study team are blinded to treatment. Treatment ends at the week-52 visit and patients return "
              "for a follow-up visit four weeks later (week 56)."),
        ("h", "4. Population"),
        ("p", "Inclusion: age 18 to 85 years; eGFR 25 to 60 mL/min/1.73 m2 at screening (central laboratory); UACR 200 to "
              "5,000 mg/g at screening; an ACE inhibitor or an angiotensin receptor blocker at the maximum tolerated dose "
              "for at least four weeks before screening. Exclusion (selected): serum potassium above 4.8 mmol/L at "
              "screening; type 1 diabetes; acute kidney injury or dialysis in the previous three months; a "
              "mineralocorticoid receptor antagonist or potassium-sparing diuretic in the previous four weeks; systolic blood "
              "pressure above 170 mmHg or below 105 mmHg."),
        ("h", "5. Study drug"),
        ("p", "One tablet is taken by mouth each morning. Bottles are dispensed at day 1 and at weeks 4, 12, 24 and 40; at "
              "a dispensing visit the patient takes that day's tablet from the new bottle after the visit's samples are "
              "taken. The first dose is taken at the day-1 visit after randomisation and after the day-1 samples; the last "
              "dose is taken on the day of the week-52 visit, after the visit's samples. Unused tablets are returned at "
              "each dispensing visit."),
        ("h", "6. Background and concomitant therapy"),
        ("p", "Investigators should keep the doses of ACE inhibitors or angiotensin receptor blockers, SGLT2 inhibitors and "
              "diuretics unchanged from randomisation to the week-12 visit unless a change is clinically necessary. From "
              "the week-12 visit, background therapy may be adjusted in line with local practice. Mineralocorticoid "
              "receptor antagonists, potassium-sparing diuretics and potassium supplements are not permitted during the "
              "study. Potassium binders may be used."),
        ("h", "7. Serum potassium"),
        ("p", "Serum potassium is measured at every visit. If serum potassium is above 6.0 mmol/L, study drug is interrupted "
              "from the next day and potassium is checked again within three to five days; study drug is restarted at the "
              "same dose once potassium is 5.0 mmol/L or below, usually about two weeks after the interruption. Study drug "
              "is stopped permanently if serum potassium is above 6.5 mmol/L, or above 6.0 mmol/L for a second time. If "
              "serum potassium is above 5.2 mmol/L and no higher than 6.0 mmol/L, the investigator repeats the test within "
              "about a week, at the central laboratory or the site's local laboratory, gives dietary advice and may start "
              "a potassium binder."),
        ("h", "8. Schedule of visits"),
        ("table", (_visit_table(), [42 * mm, 25 * mm, 18 * mm, 22 * mm, 28 * mm])),
        ("small", "Study day 1 is the day of randomisation and the first dose. Windows are counted from the target study day."),
        ("h", "9. Assessments"),
        ("p", "At every visit: serum creatinine with eGFR, potassium, sodium and glucose at the central laboratory "
              "(Tavlund Central Laboratory: the Copenhagen hub for sites in Denmark, Sweden, Germany and Poland, the Madrid "
              "hub for sites in Spain and the United Kingdom); blood pressure, pulse and weight. Urine albumin and creatinine "
              "at screening, day 1 and weeks 4, 12, 24, 52 and 56. A pre-dose sample for VMK-4410 concentration at weeks 4, "
              "12, 24 and 52. Adverse events and concomitant medications are recorded throughout."),
        ("p", "Baseline eGFR is the mean of the screening and day-1 central values; baseline UACR is the geometric mean of "
              "the screening and day-1 values."),
        ("h", "10. Stopping study drug and leaving the study"),
        ("p", "A patient may stop study drug at any time, and the investigator may stop it for safety. A patient who stops "
              "study drug permanently continues the study visits to week 56 unless consent is withdrawn. Interruptions for an "
              "intercurrent illness or a procedure are allowed; study drug is restarted when the investigator judges it safe."),
        ("h", "11. Statistics"),
        ("p", "The primary end point is analysed with a mixed model for repeated measures on log-transformed UACR. The "
              "change in eGFR at week 52 is analysed with a mixed model for repeated measures on the on-treatment values. "
              "With 150 patients per arm the study has 80% power to detect a 25% reduction in UACR against placebo at the "
              "two-sided 5% level for each dose. The statistical analysis plan gives the details."),
    ]
    _pdf(path, "VMK-4410-202 protocol version 3.0 (amendment 2)", "Clinical Development",
         "VMK-4410-202 clinical study protocol, version 3.0 (amendment 2), 10 June 2024", blocks, "2024-06-10")


# ------------------------------------------------------------------------------------------------ SAP
def sap(path):
    blocks = [
        ("title", "Statistical analysis plan, VMK-4410-202"),
        ("small", "Version 2.0, final, 28 November 2025 (version 1.0 of 2 April 2024 was written against protocol version "
                  "2.0). Varnmark Therapeutics A/S, Biostatistics. Signed before the database lock and unblinding."),
        ("h", "1. Analysis sets"),
        ("p", "Full analysis set: every randomised patient, analysed in the arm to which the patient was randomised. Safety "
              "set: every randomised patient who took at least one dose of study drug, analysed by the treatment taken."),
        ("h", "2. General definitions"),
        ("p", "Study day 1 is the date of randomisation and the first dose. The on-treatment period runs from the date of the "
              "first dose to the date of the last dose of study drug, both included; temporary interruptions of study drug "
              "fall within the on-treatment period. An on-treatment laboratory value is one from a sample taken in the "
              "on-treatment period after the first dose; the day-1 sample is taken before the first dose and is a baseline "
              "value."),
        ("p", "Baseline eGFR is the mean of the screening and day-1 central values; when one is missing, the other is used. "
              "Values are analysed under the scheduled visit under which they were collected. If a scheduled visit has more "
              "than one central value for a test, the value from the first sample is used."),
        ("h", "3. Efficacy"),
        ("p", "Primary: the change in log UACR from baseline to week 24, analysed with a mixed model for repeated measures "
              "with arm, visit, arm by visit, baseline log UACR and the randomisation strata, and an unstructured covariance; "
              "the result for each dose is given as a ratio to placebo."),
        ("p", "Secondary: the change in eGFR from baseline to week 52, analysed with the same model on the on-treatment values "
              "from weeks 2 to 52; values from samples taken after a patient has stopped study drug permanently are not "
              "used. The change from baseline to the week-56 follow-up visit is summarised by arm. eGFR analyses use the "
              "central laboratory's CKD-EPI 2021 values as reported."),
        ("h", "4. Adverse events"),
        ("p", "Adverse events are coded with MedDRA version 27.0. A treatment-emergent adverse event is one that starts on or "
              "after the date of the first dose and no later than 30 days after the date of the last dose. Adverse events of "
              "special interest: hyperkalaemia, defined by the MedDRA preferred terms Hyperkalaemia and Blood potassium "
              "increased."),
        ("p", "Partial onset dates: a missing day is taken as the first day of the month, unless the first dose was taken "
              "in that month, in which case the date of the first dose is used; a missing day and month are taken as 1 "
              "January, unless the first dose was taken in that year, in which case the date of the first dose is used. An "
              "imputed date is never later than the end date of the event."),
        ("h", "5. Laboratory safety"),
        ("p", "Serum potassium is summarised as the number of patients with an on-treatment value above 5.5 mmol/L and above "
              "6.0 mmol/L, by arm. Potassium reported in mEq/L is the same quantity as in mmol/L."),
        ("h", "6. Exposure"),
        ("p", "Duration of exposure is the number of days from the date of the first dose to the date of the last dose, both "
              "included, less the days on which study drug was interrupted. Patient-years of exposure are the sum of the "
              "durations divided by 365.25."),
        ("h", "7. Disposition"),
        ("p", "The number of patients screened, randomised, treated, completing study drug and completing the study is given "
              "by arm, with the reasons for stopping study drug and for leaving the study."),
    ]
    _pdf(path, "VMK-4410-202 statistical analysis plan version 2.0", "Biostatistics",
         "VMK-4410-202 statistical analysis plan, version 2.0, final, 28 November 2025", blocks, "2025-11-28")


# ------------------------------------------------------------------------------------------------ phase 3 synopsis
def synopsis(path):
    vs = [["Visit", "Week", "Window (days)"], ["Screening", "-2", "+/- 7"], ["Day 1 (randomisation)", "0", "none"],
          ["Weeks 2 and 4", "2, 4", "+/- 3"], ["Week 12", "12", "+/- 5"], ["Week 24", "24", "+/- 7"],
          ["Every 12 weeks from week 36 to week 144", "36 to 144", "+/- 10"], ["Month 36", "156", "+/- 10"],
          ["Off-treatment follow-up", "160", "+/- 7"]]
    blocks = [
        ("title", "VMK-4410-301 synopsis"),
        ("p", "<b>A randomised, double-blind, placebo-controlled phase 3 trial of VMK-4410 in adults with chronic kidney "
              "disease and albuminuria</b>"),
        ("small", "Draft version 0.3, 5 October 2026, for the end-of-phase-2 meeting package. Varnmark Therapeutics A/S, "
                  "Clinical Development. Not for distribution outside the study team. Changes from version 0.2: the primary "
                  "estimand written out in full; the dose-selection paragraph added."),
        ("h", "Objective"),
        ("p", "To show that VMK-4410 slows the loss of kidney function compared with placebo in adults with chronic kidney "
              "disease and albuminuria who are taking an ACE inhibitor or an angiotensin receptor blocker."),
        ("h", "Dose"),
        ("p", "One VMK-4410 dose, taken once daily, against placebo. The dose carried into this trial is the VMK-4410 dose "
              "from VMK-4410-202 expected to give the largest difference from placebo in this trial's primary end point, the "
              "change in eGFR from baseline to the month-36 visit."),
        ("h", "Design"),
        ("p", "Randomisation 1:1 to VMK-4410 or placebo, stratified as in VMK-4410-202, with 36 months of treatment for every "
              "patient. The sample size will be set once the dose and its expected difference from placebo are known."),
        ("h", "Population"),
        ("p", "As in VMK-4410-202: age 18 to 85 years; eGFR 25 to 60 mL/min/1.73 m2 and UACR 200 to 5,000 mg/g at "
              "screening; an ACE inhibitor or angiotensin receptor blocker at the maximum tolerated dose; serum potassium "
              "4.8 mmol/L or below at screening."),
        ("h", "Background therapy"),
        ("p", "As in VMK-4410-202: doses of ACE inhibitors or angiotensin receptor blockers, SGLT2 inhibitors and diuretics "
              "are kept unchanged from randomisation to the week-12 visit unless a change is clinically necessary; from the "
              "week-12 visit, background therapy may be adjusted in line with local practice. The potassium rules of "
              "VMK-4410-202 apply."),
        ("h", "Primary end point and estimand"),
        ("p", "The change in eGFR from baseline to the month-36 visit. Primary estimand: the difference between VMK-4410 and "
              "placebo in the mean change in eGFR from baseline to the month-36 visit (week 156), in all randomised "
              "patients as randomised, as if each patient had stayed on the assigned treatment: values from samples taken "
              "after a patient has stopped study drug permanently are set aside, and values during a temporary interruption "
              "are used. eGFR is calculated by the central laboratory with the CKD-EPI 2021 creatinine equation and reported "
              "with each serum creatinine. Baseline is the mean of the screening and day-1 values. The month-36 visit takes "
              "place while patients are taking study drug; the last dose is taken on the day of that visit."),
        ("h", "Secondary end points"),
        ("p", "Time to the first of a sustained fall in eGFR of 40% or more from baseline, kidney failure or death from kidney "
              "causes; the change in UACR from baseline to month 12; the change in eGFR from baseline to the off-treatment "
              "follow-up visit."),
        ("h", "Schedule of visits"),
        ("table", (vs, [70 * mm, 30 * mm, 30 * mm])),
        ("h", "Safety"),
        ("p", "Adverse events, serum potassium and serum sodium at every visit, blood pressure."),
    ]
    _pdf(path, "VMK-4410-301 synopsis, draft 0.3", "Clinical Development",
         "VMK-4410-301 synopsis, draft version 0.3, 5 October 2026", blocks, "2026-10-05")


# ------------------------------------------------------------------------------------------------ workbooks
def _hdr(ws, cols, row=1):
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE4EE")
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _widths(ws, widths):
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w


def _save(wb, path, created, modified):
    wb.save(path)
    postprocess.scrub(path, ORG, modified[:10])
    postprocess.finish_office(path, created, modified, ORG)


def lab_manual(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cover"
    for r, (a, b) in enumerate([("Tavlund Central Laboratory", "Laboratory manual for study VMK-4410-202"),
                                ("Version", "4.0"), ("Effective", "2024-06-17"), ("Sponsor", ORG),
                                ("Hubs", "Copenhagen (sites in Denmark, Sweden, Germany, Poland); Madrid (sites in Spain, United Kingdom)"),
                                ("Contact", "vmk4410-202@tavlund-lab.example")], start=1):
        ws.cell(row=r, column=1, value=a).font = Font(bold=(r == 1))
        ws.cell(row=r, column=2, value=b)
    _widths(ws, [28, 90])
    ws = wb.create_sheet("Tests")
    cols = ["Test code", "Test", "Specimen", "Copenhagen hub unit", "Madrid hub unit", "Reference range (SI)", "Method"]
    _hdr(ws, cols)
    rows = [("CREAT", "Creatinine", "Serum", "umol/L, whole number", "mg/dL, two decimals", "M 64 to 104 umol/L; F 49 to 90 umol/L",
             "Enzymatic, IDMS-traceable"),
            ("EGFR", "Glomerular filtration rate, estimated", "Serum", "mL/min/1.73m2, whole number", "mL/min/1.73m2, whole number",
             "60 and above", "CKD-EPI 2021 creatinine equation, see sheet eGFR"),
            ("K", "Potassium", "Serum", "mmol/L, one decimal", "mmol/L, one decimal", "3.5 to 5.1 mmol/L", "Indirect ion-selective electrode"),
            ("SODIUM", "Sodium", "Serum", "mmol/L", "mmol/L", "136 to 145 mmol/L", "Indirect ion-selective electrode"),
            ("GLUC", "Glucose", "Serum", "mmol/L, one decimal", "mg/dL, whole number", "3.9 to 7.8 mmol/L (non-fasting)", "Hexokinase"),
            ("MALB", "Albumin", "Urine", "mg/L", "mg/L", "below 20 mg/L", "Immunoturbidimetric"),
            ("CREATU", "Creatinine", "Urine", "mmol/L", "mg/dL", "", "Enzymatic"),
            ("ALBCREAT", "Albumin/creatinine ratio", "Urine", "mg/mmol, one decimal", "mg/g, whole number", "below 30 mg/g (3 mg/mmol)",
             "Calculated from MALB and CREATU")]
    for r, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=r, column=j, value=v)
    _widths(ws, [11, 34, 10, 24, 22, 34, 40])
    ws.freeze_panes = "A2"
    ws = wb.create_sheet("eGFR")
    lines = ["eGFR is calculated by the laboratory for every serum creatinine and reported with it.",
             "Equation: CKD-EPI 2021 creatinine equation (without a race term).",
             "eGFR = 142 x min(Scr/k, 1)^a x max(Scr/k, 1)^-1.200 x 0.9938^Age x 1.012 (if female)",
             "Scr: serum creatinine in mg/dL; k = 0.7 (female) or 0.9 (male); a = -0.241 (female) or -0.302 (male); Age in whole years on the collection date.",
             "At the Copenhagen hub creatinine is measured and reported in umol/L and converted at 88.42 umol/L per mg/dL for the calculation.",
             "At the Madrid hub creatinine is reported in mg/dL to two decimals and the calculation uses the reported value.",
             "eGFR is reported to the nearest whole number."]
    for r, t in enumerate(lines, start=1):
        ws.cell(row=r, column=1, value=t)
    _widths(ws, [130])
    ws = wb.create_sheet("Samples")
    _hdr(ws, ["Topic", "Instruction"])
    rows = [("Collection", "Blood is drawn before the day's tablet at every visit. Serum tubes are centrifuged within 60 minutes and shipped ambient on the day of collection."),
            ("Haemolysis", "Potassium is not reported on a haemolysed sample (result status NOT DONE, reason SAMPLE HAEMOLYSED, specimen condition HAEMOLYSED on the potassium result); the sample's other tests are not affected by haemolysis at the levels seen and are reported. The site is asked to collect a new sample."),
            ("Repeat samples", "A new sample collected after a haemolysed or lost sample is reported under the visit the site gives on the requisition."),
            ("Unscheduled samples", "Samples for a potassium check between visits are sent as UNSCHEDULED requisitions."),
            ("Local laboratories", "Results from a site's own laboratory are entered by the site on the local laboratory form and do not pass through the central laboratory.")]
    for r, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    _widths(ws, [22, 110])
    ws = wb.create_sheet("Changes")
    _hdr(ws, ["Version", "Effective", "Change"])
    for r, row in enumerate([("1.0", "2023-12-11", "First issue"), ("2.0", "2024-01-29", "Copenhagen hub reporting in umol/L confirmed"),
                             ("3.0", "2024-02-26", "Madrid hub added for sites in Spain and the United Kingdom"),
                             ("4.0", "2024-06-17", "Urine collection at week 4; unscheduled potassium requisitions")], start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=r, column=j, value=v)
    _widths(ws, [10, 12, 70])
    _save(wb, path, "2023-12-11T10:12:00Z", "2024-06-17T14:05:00Z")


DICT = {
    "dm.csv": [("STUDYID", "Study identifier"), ("USUBJID", "Unique subject identifier (study and subject number)"),
               ("SUBJID", "Subject number given at screening: site number and sequence"), ("SITEID", "Site number"),
               ("COUNTRY", "Country (ISO 3166 alpha-3)"), ("BRTHYR", "Year of birth"), ("AGE", "Age in years at informed consent"),
               ("AGEU", "Age unit"), ("SEX", "Sex (M, F)"), ("RACE", "Race as reported"), ("ETHNIC", "Ethnicity as reported"),
               ("RFICDTC", "Date of informed consent"), ("ARMCD", "Randomised arm code"), ("ARM", "Randomised arm"),
               ("ACTARMCD", "Arm code of the treatment taken"), ("ACTARM", "Treatment taken"),
               ("ARMNRS", "Reason the subject has no arm (screening failures)"), ("RANDDTC", "Date of randomisation"),
               ("RFXSTDTC", "Date of the first dose of study drug"), ("RFXENDTC", "Date of the last dose of study drug"),
               ("STRATEGFR", "Randomisation stratum: screening eGFR"), ("STRATUACR", "Randomisation stratum: screening UACR")],
    "sv.csv": [("STUDYID", "Study identifier"), ("USUBJID", "Unique subject identifier"), ("VISITNUM", "Visit number; unscheduled visits are numbered after the scheduled visit they follow"),
               ("VISIT", "Visit name"), ("SVSTDTC", "Visit date"), ("SVENDTC", "Visit end date"), ("SVUPDES", "Reason for an unscheduled visit")],
    "ex.csv": [("STUDYID", "Study identifier"), ("USUBJID", "Unique subject identifier"), ("EXSEQ", "Sequence number"),
               ("EXTRT", "Treatment (after unblinding)"), ("EXDOSE", "Dose per administration"), ("EXDOSU", "Dose unit"),
               ("EXDOSFRM", "Dose form"), ("EXDOSFRQ", "Dosing frequency"), ("EXROUTE", "Route"),
               ("VISIT", "Visit at which the bottles for this record were dispensed"), ("EXSTDTC", "Start date of the dosing record"),
               ("EXENDTC", "End date of the dosing record (inclusive)")],
    "ds.csv": [("STUDYID", "Study identifier"), ("USUBJID", "Unique subject identifier"), ("DSSEQ", "Sequence number"),
               ("DSTERM", "Reported term"), ("DSDECOD", "Standardised disposition term"), ("DSCAT", "Category (protocol milestone or disposition event)"),
               ("DSSCAT", "Subcategory: STUDY DRUG (end-of-treatment form) or STUDY PARTICIPATION (end-of-study form)"),
               ("EPOCH", "Study period"), ("DSSTDTC", "Date of the event")],
    "lb.csv": [("STUDYID", "Study identifier"), ("USUBJID", "Unique subject identifier"), ("LBSEQ", "Sequence number"),
               ("LBREFID", "Central laboratory accession number of the sample (blank for local laboratories)"),
               ("LBTESTCD", "Test code"), ("LBTEST", "Test name"), ("LBCAT", "Category"), ("LBSPEC", "Specimen"),
               ("LBORRES", "Result as reported"), ("LBORRESU", "Unit as reported"), ("LBSTRESC", "Result in the standard unit, as text"),
               ("LBSTRESN", "Result in the standard unit, numeric"), ("LBSTRESU", "Standard unit"),
               ("LBSTNRLO", "Reference range lower limit, standard unit"), ("LBSTNRHI", "Reference range upper limit, standard unit"),
               ("LBNRIND", "Result against the reporting laboratory's reference range"), ("LBSTAT", "NOT DONE when no result was reported"),
               ("LBREASND", "Reason no result was reported"), ("LBSPCCND", "Specimen condition"),
               ("LBMETHOD", "Method; for eGFR, the equation used"), ("LBNAM", "Laboratory that reported the result"),
               ("VISITNUM", "Visit number"), ("VISIT", "Visit name given on the requisition or the local laboratory form"),
               ("LBDTC", "Date and time of collection")],
    "ae.csv": [("STUDYID", "Study identifier"), ("USUBJID", "Unique subject identifier"), ("AESEQ", "Sequence number"),
               ("AEGRPID", "Group identifier: ties together the records of a subject that belong to one adverse event"),
               ("AETERM", "Reported term"), ("AEDECOD", "MedDRA preferred term (version 27.0)"), ("AEBODSYS", "MedDRA system organ class"),
               ("AESEV", "Severity"), ("AESER", "Serious (Y, N)"), ("AEREL", "Relationship to study drug (investigator)"),
               ("AEACN", "Action taken with study drug"), ("AEOUT", "Outcome"), ("AESTDTC", "Start date (may be partial)"),
               ("AEENDTC", "End date (blank if not resolved at the end of the study)")],
    "cm.csv": [("STUDYID", "Study identifier"), ("USUBJID", "Unique subject identifier"), ("CMSEQ", "Sequence number"),
               ("CMTRT", "Medication as reported"), ("CMDECOD", "Standardised medication name (WHODrug Global)"),
               ("CMCLAS", "ATC class (fourth level)"), ("CMCLASCD", "ATC code (fourth level)"), ("CMINDC", "Indication"),
               ("CMDOSE", "Dose per administration"), ("CMDOSTXT", "Dose as free text when the site did not enter one dose"),
               ("CMDOSU", "Dose unit"), ("CMDOSFRQ", "Frequency"), ("CMROUTE", "Route"), ("CMSTDTC", "Start date (may be partial)"),
               ("CMENDTC", "End date"), ("CMENRF", "ONGOING when the medication continued past the end of the study")],
    "vs.csv": [("STUDYID", "Study identifier"), ("USUBJID", "Unique subject identifier"), ("VSSEQ", "Sequence number"),
               ("VSTESTCD", "Test code"), ("VSTEST", "Test name"), ("VSPOS", "Position"), ("VSORRES", "Result as recorded"),
               ("VSORRESU", "Unit as recorded"), ("VSSTRESN", "Result, numeric"), ("VSSTRESU", "Unit"), ("VISITNUM", "Visit number"),
               ("VISIT", "Visit name"), ("VSDTC", "Date of measurement")],
    "pc.csv": [("STUDYID", "Study identifier"), ("USUBJID", "Unique subject identifier"), ("PCSEQ", "Sequence number"),
               ("PCTESTCD", "Analyte code"), ("PCTEST", "Analyte"), ("PCSPEC", "Specimen"), ("PCTPT", "Time point"),
               ("PCORRES", "Result as reported; below the lower limit of quantification as <0.50"), ("PCORRESU", "Unit"),
               ("PCSTRESN", "Result, numeric"), ("PCSTRESU", "Unit"), ("PCLLOQ", "Lower limit of quantification (ng/mL)"),
               ("PCSTAT", "NOT DONE when no sample was analysed"), ("VISIT", "Visit name"), ("PCDTC", "Date of collection")],
}


def data_dictionary(path, files):
    wb = Workbook()
    ws = wb.active
    ws.title = "Files"
    _hdr(ws, ["File", "Holds", "One row is"])
    holds = {"dm.csv": ("Demographics of every screened subject", "a subject"),
             "sv.csv": ("Visits that took place, scheduled and unscheduled", "a visit"),
             "ex.csv": ("Study drug taken", "a dosing record"),
             "ds.csv": ("Protocol milestones and disposition events", "an event"),
             "lb.csv": ("Laboratory results from the central laboratory and the sites' local laboratories", "a test result"),
             "ae.csv": ("Adverse events", "an adverse event record"),
             "cm.csv": ("Prior and concomitant medications", "a medication record"),
             "vs.csv": ("Vital signs", "a measurement"),
             "pc.csv": ("Trough concentrations of VMK-4410", "a sample")}
    for r, f in enumerate(files, start=2):
        ws.cell(row=r, column=1, value=f)
        ws.cell(row=r, column=2, value=holds[f][0])
        ws.cell(row=r, column=3, value=holds[f][1])
    _widths(ws, [12, 80, 22])
    for f in files:
        ws = wb.create_sheet(f.replace(".csv", ""))
        _hdr(ws, ["Column", "Description"])
        for r, (c, dsc) in enumerate(DICT[f], start=2):
            ws.cell(row=r, column=1, value=c)
            ws.cell(row=r, column=2, value=dsc)
        _widths(ws, [14, 100])
    ws = wb.create_sheet("Conventions")
    _hdr(ws, ["Topic", "Convention"])
    rows = [("Dates", "ISO 8601. Partial dates keep only the parts recorded (year, or year and month)."),
            ("Subjects", "Screening failures appear in dm.csv, ds.csv, sv.csv, lb.csv and vs.csv with their screening data only."),
            ("Unblinding", "Treatment columns are filled after the database lock of 14 April 2026."),
            ("Units", "Results keep the unit the laboratory reported in LBORRESU; LBSTRESN and LBSTRESU give every result of a test in one unit.")]
    for r, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=r, column=j, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    _widths(ws, [14, 110])
    _save(wb, path, "2026-04-15T08:02:00Z", "2026-04-15T11:47:00Z")
