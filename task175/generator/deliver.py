"""The goldens (the forecast workbook and the fiscal 2026 connections chart), generator/figures.json, submission.md and
metadata.json, all computed from the shipped files through calc.py."""
from __future__ import annotations

import json
import os

import numpy as np

import common as C
import calc

PHASES = ["Core", "1", "2", "3", "4"]
PHASE_LABEL = {"Core": "Core service area", "1": "Phase 1", "2": "Phase 2", "3": "Phase 3", "4": "Phase 4"}
QLABEL = {1: "Q1 Oct-Dec 2025", 2: "Q2 Jan-Mar 2026", 3: "Q3 Apr-Jun 2026", 4: "Q4 Jul-Sep 2026"}
COLORS = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4"]


def r5(x):
    return int(5 * round(x / 5.0))


def r50(x):
    return int(50 * round(x / 50.0))


def r10(x):
    return int(10 * round(x / 10.0))


def compute(target):
    P = calc.Pack(target)
    season = calc.season_shares(P)
    g, info = calc.forecast(P, season)
    fl = info["fl"]
    n27 = info["n27"]
    out = dict(P=P, g=g, info=info)
    out["vacant"] = {nm: int((P.area[fl] == nm).sum()) for nm in calc.FORWARD}
    out["empty"] = {nm: int(((P.area[fl] == nm) & (n27 == 0)).sum()) for nm in calc.FORWARD}
    routes = calc.connections_fy26(target, P)
    fees, lines, by_no = calc.fees_fy26(target)
    out.update(grid=routes["golden"], routes=routes, fees=fees, lines=lines, by_no=by_no)
    return out


def figures(o):
    g = o["g"]
    f = {"call_total_permits_2027": r50(g["total"])}
    f["cc1_vacant_lots_total"] = sum(o["vacant"].values())
    for nm in calc.FORWARD:
        f["cc1_vacant_lots_" + nm.lower().replace(" ", "")] = o["vacant"][nm]
    f["cc2_vacant_lots_subset_total"] = sum(o["empty"].values())
    for ph in PHASES:
        for q in (1, 2, 3, 4):
            f["ask_connections_fy2026_%s_q%d" % (ph.lower() if ph == "Core" else "phase" + ph, q)] = int(o["grid"][(ph, q)])
    for q in (1, 2, 3, 4):
        f["ask_capacity_fees_fy2026_q%d" % q] = int(round(o["fees"][q]))
    # numbers inside the area names (North 6, North 7, Northwest 3), which the write-up prints beside figures
    f["label_area_name_numbers"] = [6, 7, 3]
    f["label_call_rounding"] = 50
    # unrounded values for the bin checks
    f["_unrounded"] = {k: round(v, 3) for k, v in g.items()}
    return f


def write_workbook(path, o, f):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    import postprocess as PP
    wb = Workbook()
    ws = wb.active
    ws.title = "2027 permit forecast"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDE7F3")
    ws["A1"] = "Single-family permits in the 2027 program areas: %d in calendar 2027" % f["call_total_permits_2027"]
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Port Ardell Utilities, Finance Division, Revenue and Rates. Working file for the January 2027 revenue "
                "sufficiency report. Records to 31 December 2026; building market held at its calendar 2026 level.")
    ws["A2"].font = Font(italic=True, size=9)
    hdr = ["Area", "Plat units", "Lines accepted", "Vacant lots, 1 Jan 2027", "On streets with no home standing"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = bold; c.fill = head_fill; c.alignment = Alignment(wrap_text=True, vertical="top")
    P = o["P"]
    r = 5
    for nm in calc.FORWARD:
        a = P.areas[nm]
        ws.cell(row=r, column=1, value=nm)
        ws.cell(row=r, column=2, value=", ".join(str(u) for u in a["units"]))
        c = ws.cell(row=r, column=3, value=a["accepted"]); c.number_format = "d mmm yyyy"
        c = ws.cell(row=r, column=4, value=o["vacant"][nm]); c.number_format = "#,##0"
        c = ws.cell(row=r, column=5, value=o["empty"][nm]); c.number_format = "#,##0"
        r += 1
    ws.cell(row=r, column=1, value="Three areas").font = bold
    c = ws.cell(row=r, column=4, value=sum(o["vacant"].values())); c.number_format = "#,##0"; c.font = bold
    c = ws.cell(row=r, column=5, value=sum(o["empty"].values())); c.number_format = "#,##0"; c.font = bold
    top = Border(top=Side(style="thin"))
    for j in range(1, 6):
        ws.cell(row=r, column=j).border = top
    ws.cell(row=r + 2, column=1, value="Single-family permits forecast for 2027, the three areas together (nearest 50)").font = bold
    c = ws.cell(row=r + 2, column=5, value=f["call_total_permits_2027"]); c.number_format = "#,##0"; c.font = bold
    ws.cell(row=r + 4, column=1, value=("The forecast is the expected number of distinct lots permitted in 2027. Each "
                                        "vacant lot's 2027 chance of a permit is taken from "
                                        "served vacant lots 2008 to 2026 with as many homes standing on their street (both "
                                        "sides of the hundred-block) on 1 January, canal lots apart, at the 2026 level; "
                                        "Northwest 3 lots count at the well and septic rate until its lines are accepted.")).font = Font(size=9)
    ws.cell(row=r + 5, column=1, value="Source: Calusa County Property Appraiser 2026 roll; Building Division permits 2008 to 2026; "
                                       "program schedule Revision 12.").font = Font(size=9, italic=True)
    for col, wv in zip("ABCDE", [16, 13, 14, 14, 17]):
        ws.column_dimensions[col].width = wv
    ws.row_dimensions[4].height = 45
    ws.freeze_panes = "A5"

    ws2 = wb.create_sheet("FY2026 connections")
    ws2["A1"] = "New single-family connections, fiscal 2026"
    ws2["A1"].font = Font(bold=True, size=13)
    ws2["A2"] = ("A lot's first meter install, at a lot with no account opened before it, premises placed on lots by "
                 "parcel or service address; quarter of the install date.")
    ws2["A2"].font = Font(italic=True, size=9)
    hdr2 = ["Phase"] + [QLABEL[q] for q in (1, 2, 3, 4)] + ["Fiscal 2026"]
    for j, h in enumerate(hdr2, 1):
        c = ws2.cell(row=4, column=j, value=h); c.font = bold; c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
    r = 5
    for ph in PHASES:
        ws2.cell(row=r, column=1, value=PHASE_LABEL[ph])
        for q in (1, 2, 3, 4):
            c = ws2.cell(row=r, column=1 + q, value=int(o["grid"][(ph, q)])); c.number_format = "#,##0"
        c = ws2.cell(row=r, column=6, value=int(sum(o["grid"][(ph, q)] for q in (1, 2, 3, 4)))); c.number_format = "#,##0"
        r += 1
    ws2.cell(row=r, column=1, value="All served areas").font = bold
    for q in (1, 2, 3, 4):
        c = ws2.cell(row=r, column=1 + q, value=int(sum(o["grid"][(ph, q)] for ph in PHASES))); c.number_format = "#,##0"; c.font = bold
    c = ws2.cell(row=r, column=6, value=int(sum(o["grid"].values()))); c.number_format = "#,##0"; c.font = bold
    for j in range(1, 7):
        ws2.cell(row=r, column=j).border = top
    ws2.cell(row=r + 2, column=1, value=("AMI head-end installs at lots already billed before the install are meter "
                                         "exchanges, and a later install at a lot, on its stub or on a premise set up at the counter, is a meter set again, not a connection. "
                                         "Phases as the program schedule groups the areas.")).font = Font(size=9)
    for col, wv in zip("ABCDEF", [20, 13, 13, 13, 13, 12]):
        ws2.column_dimensions[col].width = wv
    ws2.row_dimensions[4].height = 30
    ws2.freeze_panes = "B5"

    ws3 = wb.create_sheet("FY2026 capacity fees")
    ws3["A1"] = "Water and sewer capacity fees received, fiscal 2026"
    ws3["A1"].font = Font(bold=True, size=13)
    ws3["A2"] = "Cashiering receipts, fee code CAPWS, by receipt date; each payment (receipt number and payer) counted once."
    ws3["A2"].font = Font(italic=True, size=9)
    for j, h in enumerate(["Quarter", "Receipts", "Fees received (USD)"], 1):
        c = ws3.cell(row=4, column=j, value=h); c.font = bold; c.fill = head_fill
    r = 5
    for q in (1, 2, 3, 4):
        ws3.cell(row=r, column=1, value=QLABEL[q])
        c = ws3.cell(row=r, column=3, value=f["ask_capacity_fees_fy2026_q%d" % q]); c.number_format = "#,##0"
        r += 1
    ws3.delete_cols(2)
    ws3.cell(row=4, column=2, value="Fees received (USD)").font = bold
    ws3.cell(row=4, column=2).fill = head_fill
    ws3.cell(row=r, column=1, value="Fiscal 2026").font = bold
    c = ws3.cell(row=r, column=2, value=sum(f["ask_capacity_fees_fy2026_q%d" % q] for q in (1, 2, 3, 4)))
    c.number_format = "#,##0"; c.font = bold
    for j in (1, 2):
        ws3.cell(row=r, column=j).border = top
    ws3.cell(row=r + 2, column=1, value=("A builder's receipt covering several permits carries the receipt amount on each "
                                         "permit line; receipt numbers run by drawer and day.")).font = Font(size=9)
    ws3.column_dimensions["A"].width = 20
    ws3.column_dimensions["B"].width = 20
    wb.save(path)
    PP.normalize_ooxml(path, "Melissa Payne", "Melissa Payne", "2027-01-11T15:30:00Z", "2027-01-12T21:05:00Z", excel=True)


def write_chart(path, o):
    import matplotlib
    matplotlib.use("Agg")
    matplotlib.rcParams["svg.hashsalt"] = "task175"
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(8.6, 5.4), dpi=150)
    fig.patch.set_facecolor("#fcfcfb"); ax.set_facecolor("#fcfcfb")
    qs = [1, 2, 3, 4]
    bottom = np.zeros(4)
    for ph, col in zip(PHASES, COLORS):
        vals = np.array([o["grid"][(ph, q)] for q in qs], dtype=float)
        bars = ax.bar(range(4), vals, bottom=bottom, color=col, width=0.62, edgecolor="#fcfcfb", linewidth=2,
                      label=PHASE_LABEL[ph])
        for k, (v, b) in enumerate(zip(vals, bottom)):
            ax.text(k, b + v / 2, "%d" % v, ha="center", va="center", fontsize=9, color="#0b0b0b")
        bottom += vals
    for k, tot in enumerate(bottom):
        ax.text(k, tot + 6, "%d" % tot, ha="center", va="bottom", fontsize=9.5, color="#52514e", fontweight="bold")
    ax.set_xticks(range(4))
    ax.set_xticklabels([QLABEL[q].replace(" ", "\n", 1) for q in qs], fontsize=9, color="#52514e")
    ax.set_ylabel("New single-family connections", fontsize=9.5, color="#52514e")
    ax.set_title("Fiscal 2026 new single-family connections by quarter: {:,} in all".format(int(bottom.sum())), fontsize=11.5,
                 loc="left", color="#0b0b0b")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.spines["left"].set_color("#b9b8b3"); ax.spines["bottom"].set_color("#b9b8b3")
    ax.yaxis.grid(True, color="#e8e7e3", linewidth=0.8); ax.set_axisbelow(True)
    ax.tick_params(colors="#52514e", labelsize=9)
    ax.set_ylim(0, bottom.max() * 1.12)
    ax.legend(loc="upper left", bbox_to_anchor=(1.0, 1.0), frameon=False, fontsize=9)
    fig.text(0.01, 0.01, "Source: AMI head-end installs and CIS account history; first install per lot only, meter exchanges at billed homes excluded. "
             "Port Ardell Utilities, Revenue and Rates.", fontsize=7.5, color="#52514e")
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    fig.savefig(path, facecolor=fig.get_facecolor(), metadata={"Software": None})
    plt.close(fig)


def write_submission(path, o, f):
    g = o["g"]
    lines = []
    lines.append("## Tags\n")
    lines.append("**Domain:** Urban and regional science, land use and spatial development modeling.  ")
    lines.append("**Analytical objective:** Forecasting and predictive modeling.\n")
    lines.append("## 1. Final Recommendation\n")
    lines.append("**Port Ardell will issue about %d single-family building permits in 2027 on lots in North 6, North 7 "
                 "and Northwest 3 together, to the nearest 50.**" % f["call_total_permits_2027"])
    lines.append("The figure takes each vacant lot's chance of a permit from served vacant lots on streets with as many "
                 "homes standing as its own, and %s of the %s vacant lots stand on streets with no home. Not the "
                 "first-year permits per vacant lot of the earlier extension areas, which were connected when most of "
                 "their streets already had homes. Not a lot model without the street, or one built on the plat block, "
                 "which mixes the lots of two streets.\n" % ("{:,}".format(f["cc2_vacant_lots_subset_total"]),
                                                             "{:,}".format(f["cc1_vacant_lots_total"])))
    lines.append("## 2. Critical Components\n")
    lines.append("1. Vacant lots on 1 January 2027: **%s** in the three areas (North 6 **%s**, North 7 **%s**, Northwest 3 **%s**)" % tuple(
        "{:,}".format(x) for x in (f["cc1_vacant_lots_total"], f["cc1_vacant_lots_north6"],
                                   f["cc1_vacant_lots_north7"], f["cc1_vacant_lots_northwest3"])))
    lines.append("2. Vacant lots on streets with no home standing on 1 January 2027: **%s**" %
                 "{:,}".format(f["cc2_vacant_lots_subset_total"]))
    lines.append("")
    lines.append("## 3. Step-by-Step Solution\n")
    steps = [
        "Placed every lot in its area through the plat unit in its legal description in pa_parcel_extract_2026.csv and "
        "the plat units and acceptance dates in uep_program_schedule_rev12_2026.pdf.",
        "Joined the Residential New Single Family permits of bldg_permits_legacy_2008_2018.csv (by site address) and "
        "bldg_permits_2019_2026.xlsx (by parcel), dropping the legacy copy of each carried-over permit, for each lot's "
        "issue and certificate of occupancy dates.",
        "Counted the homes standing on each lot's street (the street and hundred-block of its site address, both sides) on "
        "1 January of each year from YR_BUILT and the certificates of occupancy issued in 2026, which the 2026 roll does "
        "not yet show.",
        "Fitted served vacant lot-years 2008 to 2026 (exposure from each area's acceptance date to the permit) as a "
        "Poisson rate: a year effect times a canal effect times a separate effect for each band of homes standing on "
        "the street (0, 1, 2, 3, 4 to 5, 6 to 8, 9 to 12, 13 to 17, 18 or more), not a slope on the count.",
        "Applied the 2026 level, as the forecasting conventions require, to the %s vacant lots of the three areas by "
        "their streets and canals, each area served from its scheduled acceptance date and at the well and septic rate "
        "before it, and summed each lot's chance of a permit in 2027 (one less the exponential of its cumulative rate, "
        "the expected number of distinct lots the conventions define): %d permits." % (
            "{:,}".format(f["cc1_vacant_lots_total"]), f["call_total_permits_2027"]),
        "Placed every premise in cis_premise_extract_20261231.csv on its lot by its parcel number where recorded and "
        "otherwise by service address (street types normalised), and took fiscal 2026 new connections from ami_headend_installs_fy25_fy26.csv as each lot's "
        "first install in the log, when it falls in fiscal 2026 at a lot with no account in "
        "cis_account_history_20261231.csv opened before it: a lot carried under two premises (its stub and one set up "
        "at the counter when a stalled build's meter was set again) is one connection whatever their address strings, a later install at a lot is a "
        "meter set again, and an install at a billed lot is a meter exchange; phases through the plat unit into the "
        "schedule's phases rather than by the premise's stored UEP_PHASE.",
        "Summed capacity_fee_receipts_fy2026.xlsx once per payment by receipt date, a payment being a receipt number "
        "and its payer: a receipt covering several permits repeats its amount on each permit line, and receipt "
        "numbers run by drawer and day, so two payers' receipts can share a number.",
    ]
    for k, s in enumerate(steps, 1):
        lines.append("%d. %s" % (k, s))
    lines.append("")
    lines.append("## 4. Deliverable Answers\n")
    lines.append("### %s\n" % C.DELIV_XLSX)
    lines.append("1. Single-family building permits in 2027 in the three areas together: **%d** (nearest 50)" % f["call_total_permits_2027"])
    lines.append("2. New single-family connections in fiscal 2026, by quarter (Q1 October to December 2025 through Q4 July to September 2026):")
    for ph in PHASES:
        key = ph.lower() if ph == "Core" else "phase" + ph
        lines.append("   - %s: Q1 %d, Q2 %d, Q3 %d, Q4 %d" % ((PHASE_LABEL[ph],) + tuple(
            f["ask_connections_fy2026_%s_q%d" % (key, q)] for q in (1, 2, 3, 4))))
    lines.append("3. Water and sewer capacity fees received in fiscal 2026: Q1 **USD %s**, Q2 **USD %s**, Q3 **USD %s**, Q4 **USD %s**" % tuple(
        "{:,}".format(f["ask_capacity_fees_fy2026_q%d" % q]) for q in (1, 2, 3, 4)))
    lines.append("")
    lines.append("### %s\n" % C.DELIV_PNG)
    lines.append("1. One stacked bar per fiscal 2026 quarter with a segment for the core service area and each of Phases 1 to 4, "
                 "each segment labelled with its connection count as in answer 2 above.")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")


def write_metadata(path):
    meta = {"task": "task175", "distractor_files": ["target/" + d for d in C.DISTRACTORS],
            "deliverables": [C.DELIV_XLSX, C.DELIV_PNG], "objective": "Forecasting & Predictive Modeling",
            "domain": "Urban & Regional Science / Land Use & Spatial Development Modeling"}
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(meta, fh, indent=2)
        fh.write("\n")


def main(target, golden, figpath, subpath, metapath):
    o = compute(target)
    f = figures(o)
    os.makedirs(golden, exist_ok=True)
    write_workbook(os.path.join(golden, C.DELIV_XLSX), o, f)
    write_chart(os.path.join(golden, C.DELIV_PNG), o)
    with open(figpath, "w", encoding="utf-8") as fh:
        json.dump(f, fh, indent=2, sort_keys=True)
        fh.write("\n")
    write_submission(subpath, o, f)
    write_metadata(metapath)
    return o, f
