"""
business.py, the people and business-furniture layer of the fingerprint guard.

Two repetitions the mechanism cards cannot see. The first is people: one first name turned up in
32 builds with ten different surnames, because a model asked to invent a colleague reaches for the
same handful of names. The second is the business furniture around the decision: who decides, at
what forum, under what deadline, inside what kind of organisation. 28 prompts put the call in front
of a board.

People are detected only in the files a reader meets as narrative (the prompt, the write-up, the
pack's documents, threads and memos, the goldens), never in data tables, where generated customer
names are rows rather than personas. Detection needs a first name from Faker's locale lists
followed by a capitalised surname that is not a common word and not the start of an organisation or
place name.
"""

import importlib
import json
import os
import re
from collections import Counter, defaultdict
from pathlib import Path

NARRATIVE_EXT = {".md", ".txt", ".eml", ".html", ".htm", ".pdf", ".docx", ".pptx", ".rtf"}
TOP_LEVEL_FILES = ("prompt.md", "PROMPT.md", "problem_statement.md", "original_prompt.txt",
                   "submission.md", "submissions.md", "SOLUTION.md", "solution.md")
NARRATIVE_DIRS = ("target", "golden", "deliverables", "golden_deliverables", "answer_key")

# Capitalised words that are first names in some locale and something else far more often in a
# business document: months, places, titles, common nouns. A person whose first name is one of
# these is missed, which is the cheaper error than counting "June Forecast" as a colleague.
AMBIGUOUS_FIRST = set("""
May June April August March Mark Grant Bill Chase Page Rose Hope Faith Major Royal Summer Winter
Autumn Dawn Joy Grace Lincoln Austin Dallas Houston Jordan Paris Sydney Victoria Florence Georgia
Carolina Virginia Kent York Chester Warren Marshall Sherman Stanley Dean Bishop Judge Christian
Hunter Carter Mason Taylor Parker Porter Cooper Baker Fisher Miles Frank Price Lane Hill Dale Glen
Brook Ray Star Sterling Penny Ruby Crystal Amber Ivy Holly Iris Lily Daisy Violet Heather Olive
Jasmine Sandy Sunny Wade Drew Gene Earl Duke Prince King Queen Baron Lord Saint Justice Liberty
Unity Harmony Haven Honor Merit Noble Reed Rowan Sage Scout Shelby Spring Storm True Tyler Vale
Wren Young West North South East Long Case Hall Marsh Moss Park Read Ross Shaw Ward Wells Wood Bell
Bush Field Fox Gold Green Gray Grey Ford Clay Cash Cruz Sol Van Von Del Dee Kai Mac Will Rich Pat
Don Art Bob Sue Gay Coy Buck Bud Colt Crew Fleet Payne Tom Many Monday Friday Sunday Total Cost
Rate Delta Alpha Beta Gamma Sierra Nevada Montana Dakota Cheyenne Phoenix Denver Aspen Easton
Madison Jackson Clinton Hamilton Franklin Jefferson Washington Adams Monroe Harrison Wilson Nelson
Lewis Allen Scott Morgan Kelly Kerry Shannon Logan Regan Quinn Avery Blair Brooke Carson Dakota
Devon Emery Finley Harley Hayden Jamie Kendall Kennedy Lindsay Marley Peyton Reagan Riley Sawyer
Sydney Tatum Whitney Ashley Sidney Meridian Summit Harbor Harbour Valley River Lake Bay Sterling
""".split())

# Words that cannot be a surname in these documents. The second token of a candidate name is
# rejected when it is one of these, and the pair is rejected when the word after it is one.
NOT_SURNAME = set("""
Street Road Avenue Way Drive Lane Court Place Square Terrace Gardens Close Crescent Parade Row
Council County City Town Village Borough District Region Regional Province State Federal National
Foundation Fund Trust Bank Group Partners Partnership Holdings Health Hospital Clinic Medical Care
School Schools College University Institute Centre Center Academy Library Museum Church Office
Department Agency Authority Board Commission Committee Company Corporation Inc Ltd Limited LLC PLC
Services Service Systems Solutions Logistics Industries Manufacturing Valley River Bay Harbour
Harbor Park Hills Heights Central Station Depot Warehouse Plant Mill Works Line Lines Freight
Transport Rail Railway Airport Port Market Mall Store Stores Shop Hall Housing Homes Estates
Foodservice Foods Food Farms Farm Energy Power Water Utilities Utility Network Networks Capital
Credit Union Insurance Mutual Cooperative Coop Co Association Alliance Society Network Labs Lab
Technologies Technology Software Analytics Data Media Press News Studio Studios Global
International Pacific Atlantic Northern Southern Eastern Western Mountain Coast Coastal Island
Islands Springs Falls Creek Ridge Forest Woods Point Beach Bridge Gate Cross Junction Crossing
Programme Program Project Plan Policy Standard Rule Rules Act Code Manual Guide Report Review
Summary Notes Note Memo Brief Thread Minutes Agenda Register Ledger Schedule Table Index Annex
Appendix Section Chapter Part Item Figure Chart Total Year Quarter Month Week Day Fiscal Budget
""".split())

CALENDAR = set("""January February March April May June July August September October November
December Monday Tuesday Wednesday Thursday Friday Saturday Sunday""".split())
NOT_SURNAME |= CALENDAR
AMBIGUOUS_FIRST |= CALENDAR | {"Golden", "General", "Council"}


def _dictionary():
    """Lower-case English words from the system word list. A capitalised token whose lower-case
    form is an ordinary word is a heading or an organisation far more often than a person."""
    try:
        with open("/usr/share/dict/words", encoding="utf-8", errors="replace") as fh:
            return {w.strip() for w in fh if w[:1].islower()}
    except Exception:
        return set()


DICT_LOWER = _dictionary()
_LAST = None


def last_names():
    global _LAST
    if _LAST is not None:
        return _LAST
    names = set()
    try:
        import faker.config
        locales = list(faker.config.AVAILABLE_LOCALES)
    except Exception:
        locales = []
    for loc in locales:
        try:
            prov = getattr(importlib.import_module("faker.providers.person.%s" % loc), "Provider", None)
        except Exception:
            continue
        val = getattr(prov, "last_names", None) if prov else None
        if not val or isinstance(val, property):
            continue
        try:
            items = list(val.keys() if hasattr(val, "keys") else val)
        except TypeError:
            continue
        names.update(str(n).strip() for n in items)
    _LAST = names
    return _LAST


LOCALE_FOR_GEO = [
    (r"\b(united kingdom|uk|england|scotland|wales|britain)\b", "en_GB"),
    (r"\b(ireland)\b", "en_IE"),
    (r"\b(australia)\b", "en_AU"),
    (r"\b(new zealand)\b", "en_NZ"),
    (r"\b(canada)\b.*\b(quebec|montreal)\b", "fr_CA"),
    (r"\b(canada)\b", "en_CA"),
    (r"\b(india)\b", "en_IN"),
    (r"\b(germany)\b", "de_DE"),
    (r"\b(austria)\b", "de_AT"),
    (r"\b(switzerland)\b", "de_CH"),
    (r"\b(france)\b", "fr_FR"),
    (r"\b(spain)\b", "es_ES"),
    (r"\b(mexico)\b", "es_MX"),
    (r"\b(italy)\b", "it_IT"),
    (r"\b(brazil)\b", "pt_BR"),
    (r"\b(portugal)\b", "pt_PT"),
    (r"\b(netherlands|holland)\b", "nl_NL"),
    (r"\b(belgium)\b", "nl_BE"),
    (r"\b(sweden)\b", "sv_SE"),
    (r"\b(norway)\b", "no_NO"),
    (r"\b(denmark)\b", "da_DK"),
    (r"\b(finland)\b", "fi_FI"),
    (r"\b(poland)\b", "pl_PL"),
    (r"\b(czech)\b", "cs_CZ"),
    (r"\b(south africa)\b", "en_US"),
    (r"\b(philippines)\b", "en_PH"),
    (r"\b(united states|usa|us)\b", "en_US"),
]


def locale_for(geography):
    g = (geography or "").lower()
    for pat, loc in LOCALE_FOR_GEO:
        if re.search(pat, g):
            return loc
    return "en_US"


# ------------------------------------------------------------------ first-name dictionary

_FIRST = None


def first_names():
    """Every first name Faker carries in a Latin-script locale, with the ambiguous ones removed."""
    global _FIRST
    if _FIRST is not None:
        return _FIRST
    names = set()
    try:
        import faker.config
        locales = list(faker.config.AVAILABLE_LOCALES)
    except Exception:
        locales = []
    for loc in locales:
        try:
            mod = importlib.import_module("faker.providers.person.%s" % loc)
        except Exception:
            continue
        prov = getattr(mod, "Provider", None)
        if prov is None:
            continue
        for attr in ("first_names", "first_names_male", "first_names_female", "first_names_nonbinary"):
            val = getattr(prov, attr, None)
            if not val or isinstance(val, property):
                continue
            try:
                items = list(val.keys() if hasattr(val, "keys") else val)
            except TypeError:
                continue
            for n in items:
                n = str(n).strip()
                if re.fullmatch(r"[A-Z][a-z\u00e0-\u00ff]{2,14}", n):
                    names.add(n)
    _FIRST = names - AMBIGUOUS_FIRST
    return _FIRST


# ------------------------------------------------------------------ detection

NAME_PAIR = re.compile(r"\b([A-Z][a-z\u00e0-\u00ff]{2,14})\s+((?:Mc|Mac|O')?[A-Z][a-z\u00e0-\u00ff]{1,20}(?:-[A-Z][a-z\u00e0-\u00ff]{2,20})?)(?:\s+([A-Z][A-Za-z&]+))?")


def people_in_text(text):
    """Person names (first, last) in a narrative text."""
    fn, ln = first_names(), last_names()
    out = set()
    for m in NAME_PAIR.finditer(text or ""):
        first, last, nxt = m.group(1), m.group(2), m.group(3)
        if first not in fn or first.lower() in DICT_LOWER or last in NOT_SURNAME:
            continue
        if last not in ln and last.lower() in DICT_LOWER:
            continue
        if nxt and nxt in NOT_SURNAME:
            continue
        if last.lower() in ("and", "the", "of", "for"):
            continue
        out.add((first, last))
    return out


def read_narrative(path, cap=300000):
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            import logging
            logging.getLogger("pypdf").setLevel(logging.ERROR)
            from pypdf import PdfReader
            r = PdfReader(str(path))
            return "\n".join((p.extract_text() or "") for p in r.pages[:40])[:cap]
        if ext == ".docx":
            import docx
            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" ".join(c.text for c in row.cells))
            return "\n".join(parts)[:cap]
        if ext == ".pptx":
            try:
                from pptx import Presentation
            except Exception:
                return ""
            out = []
            for s in Presentation(str(path)).slides:
                for sh in s.shapes:
                    if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                        out.append(sh.text_frame.text)
            return "\n".join(out)[:cap]
        with open(path, "r", encoding="utf-8", errors="replace") as fh:
            return fh.read(cap)
    except Exception:
        return ""


def narrative_files(task_dir):
    d = Path(task_dir)
    files = [d / f for f in TOP_LEVEL_FILES if (d / f).exists()]
    for sub in NARRATIVE_DIRS:
        root = d / sub
        if not root.is_dir():
            continue
        for p in sorted(root.rglob("*")):
            if p.is_file() and p.suffix.lower() in NARRATIVE_EXT and not p.name.startswith("."):
                if p.stat().st_size < 3_000_000:
                    files.append(p)
    return files


def people_in_build(task_dir):
    found = set()
    for p in narrative_files(task_dir):
        found |= people_in_text(read_narrative(p))
    return sorted(found)


# ------------------------------------------------------------------ the corpus index

def build_index(repo, cache_path):
    builds = {}
    for d in sorted(Path(repo).glob("task*")):
        if d.is_dir() and re.fullmatch(r"task\d+", d.name):
            builds[d.name] = [list(x) for x in people_in_build(d)]
    data = {"builds": builds}
    Path(cache_path).parent.mkdir(parents=True, exist_ok=True)
    # Written to a temporary file and renamed, so two guard runs refreshing at once can never leave
    # a half-written index behind.
    tmp = Path(str(cache_path) + ".%d.tmp" % os.getpid())
    tmp.write_text(json.dumps(data, indent=1, ensure_ascii=False))
    os.replace(tmp, cache_path)
    return data


def load_index(repo, cache_path, refresh=False):
    p = Path(cache_path)
    if refresh or not p.exists():
        return build_index(repo, cache_path)
    newest = max((d.stat().st_mtime for d in Path(repo).glob("task*") if d.is_dir()), default=0)
    if p.stat().st_mtime < newest:
        return build_index(repo, cache_path)
    return json.loads(p.read_text())


def usage(index, exclude=()):
    firsts, lasts, fulls = defaultdict(set), defaultdict(set), defaultdict(set)
    for task, pairs in index["builds"].items():
        if task in exclude:
            continue
        for first, last in pairs:
            firsts[first].add(task)
            lasts[last].add(task)
            fulls["%s %s" % (first, last)].add(task)
    return firsts, lasts, fulls


def check_people(people, index, recent_tasks, me):
    """Ordinary names that are not recycled. Returns (level, rule, message, against) tuples."""
    firsts, lasts, fulls = usage(index, exclude={me})
    recent = set(recent_tasks)
    out = []
    seen_f, seen_l = Counter(), Counter()
    for full in people or []:
        parts = full.split()
        if len(parts) < 2:
            out.append(("BLOCK", "people.form", "%r needs a first name and a surname" % full, []))
            continue
        first, last = parts[0], parts[-1]
        seen_f[first] += 1
        seen_l[last] += 1
        if fulls.get(full):
            out.append(("BLOCK", "people.full", "%s is already a persona elsewhere" % full, sorted(fulls[full])))
        uf = firsts.get(first, set())
        if len(uf) >= 3 or (uf & recent):
            out.append(("BLOCK", "people.first", "first name %s is used in %d builds%s" % (
                first, len(uf), ", including a recent one" if uf & recent else ""), sorted(uf)[:8]))
        elif uf:
            out.append(("WARN", "people.first", "first name %s appears in an older build" % first, sorted(uf)))
        ul = lasts.get(last, set())
        if len(ul) >= 2 or (ul & recent):
            out.append(("BLOCK", "people.last", "surname %s is used in %d builds%s" % (
                last, len(ul), ", including a recent one" if ul & recent else ""), sorted(ul)[:8]))
        elif ul:
            out.append(("WARN", "people.last", "surname %s appears in an older build" % last, sorted(ul)))
    for f, n in seen_f.items():
        if n > 1:
            out.append(("BLOCK", "people.inside", "two personas in this build share the first name %s" % f, []))
    for l, n in seen_l.items():
        if n > 1:
            out.append(("BLOCK", "people.inside", "two personas in this build share the surname %s" % l, []))
    return out


def suggest_people(index, recent_tasks, locale, n, seed):
    """Draw names from the locale's own pool, never by hand, skipping anything the check would block
    and anything already drawn in this list."""
    from faker import Faker
    fk = Faker(locale)
    fk.seed_instance(seed)
    firsts, lasts, fulls = usage(index)
    recent = set(recent_tasks)
    fn = first_names()
    out, used_f, used_l = [], set(), set()
    for _ in range(5000):
        if len(out) >= n:
            break
        first, last = fk.first_name(), fk.last_name()
        if not re.fullmatch(r"[A-Z][A-Za-z\u00e0-\u00ff'\-]{1,20}", first) or not re.fullmatch(r"[A-Z][A-Za-z\u00e0-\u00ff'\-]{1,24}", last):
            continue
        if first in AMBIGUOUS_FIRST or last in NOT_SURNAME or first in used_f or last in used_l:
            continue
        uf, ul = firsts.get(first, set()), lasts.get(last, set())
        if len(uf) >= 2 or (uf & recent) or ul or fulls.get("%s %s" % (first, last)):
            continue
        out.append("%s %s" % (first, last))
        used_f.add(first)
        used_l.add(last)
    return out


# ------------------------------------------------------------------ business furniture

FORUM_RULES = [
    ("board_of_directors", r"\b(the|our|a|my) board\b(?! game)|\bboard (meeting|meets|votes?|papers?|pack|approval|paper)\b|\bboard of directors\b"),
    ("trustees_or_governors", r"\btrustees\b|\bgovernors\b|\bboard of trustees\b"),
    ("council_or_assembly", r"\b(the|our|city|county|town|state|district) council\b|\bcouncil (meeting|votes?|sits|session)\b|\bassembly\b|\blegislature\b"),
    ("committee_or_panel", r"\bcommittee\b|\bpanel\b|\bsteering group\b|\bworking group\b"),
    ("regulator_or_inspector", r"\bregulator\b|\binspectorate\b|\binspector\b|\bombudsman\b"),
    ("minister_or_cabinet", r"\bminister\b|\bcabinet\b|\bsecretary of state\b"),
    ("executive_team", r"\b(executive|leadership|management|exec) (team|meeting|review|committee)\b|\bthe (ceo|cfo|coo)\b"),
    ("funder_or_donor", r"\b(the|our) (funder|donor|grantor)\b"),
    ("auditor", r"\b(the|our|external|internal) auditors?\b"),
    ("customer_or_counterparty", r"\b(the|our) (customer|client|carrier|supplier|vendor|insurer|lender|landlord)\b"),
]

EVENT_RULES = [
    ("vote_or_meeting", r"\b(votes?|voting) (on|at)\b|\bmeets? on\b|\b(meeting|session|sitting) (on|next|this)\b|\bagenda\b"),
    ("statutory_or_regulatory_filing", r"\b(file|files|filing|lodge|lodges|submit|submits) (the|our|a|it|with)\b|\breturn (is )?due\b|\bcertif(y|ies|ication)\b"),
    ("purchase_order_or_contract", r"\bpurchase order\b|\b(the|a) contract\b|\brenewal\b|\btender\b|\bRFP\b"),
    ("budget_or_appropriation", r"\bbudget\b|\bappropriation\b|\bfunding round\b|\ballocation round\b"),
    ("period_close", r"\b(month|quarter|year)[- ]end\b|\b(at|before) the close\b|\bclose of (the )?(month|quarter|year|books)\b"),
    ("audit_or_inspection", r"\baudit\b|\binspection\b"),
    ("cutover_or_migration", r"\bcutover\b|\bmigration\b|\bgo[- ]live\b|\bswitchover\b"),
    ("launch_or_rollout", r"\blaunch\b|\brollout\b|\broll out\b"),
    ("season_or_peak", r"\bpeak\b|\bseason\b|\bharvest\b|\bholiday rush\b"),
    ("incident_or_complaint", r"\bincident\b|\bcomplaints?\b|\boutage\b|\brecall\b"),
]

ORG_RULES = [
    ("foundation_or_funder", r"foundation|funder|philanthrop|endowment|grantmak"),
    ("local_government", r"\bcounty\b|\bcity\b|\bmunicipal|\bcouncil\b|\bborough\b|\btown\b|\bparish\b"),
    ("government_agency", r"\bstate\b|federal|ministry|department|\bagency\b|government|public (body|authority)|\boffice\b|revenue|treasury|central bank|authority"),
    ("university_or_school", r"universit|college|school|district|academy|education"),
    ("hospital_or_health_provider", r"hospital|health|clinic|care provider|medical|pharma|nhs"),
    ("financial_institution", r"\bbank\b|credit union|lender|insur|fintech|payments?|card issuer|asset manag"),
    ("utility_or_infrastructure", r"utilit|water|power|energy|transit|rail|airport|port authority|telecom"),
    ("logistics_or_3pl", r"logistic|3pl|carrier|freight|distribut|warehous|courier|haulier|shipping"),
    ("manufacturer", r"manufactur|plant|factory|fabricat|mill|assembl|industrial"),
    ("retailer_or_ecommerce", r"retail|store|e-?commerce|marketplace|grocer|shop"),
    ("software_or_platform", r"software|platform|saas|\bapp\b|subscription|streaming|marketplace app|developer"),
    ("nonprofit_or_charity", r"nonprofit|non-profit|charit|community|ngo|food bank|shelter|association"),
    ("research_or_statistics_office", r"statistic|bureau|research|survey|census"),
    ("cooperative_or_member_body", r"cooperative|co-op|mutual|member|union|trade body|chamber"),
]


def classify(text, rules):
    t = (text or "").lower()
    return [name for name, pat in rules if re.search(pat, t, re.I)]


def derived_elements(repo, task):
    """Forum and forcing-event sets read off a build's prompt, for cards written before these
    fields existed."""
    d = Path(repo) / task
    text = ""
    for f in ("prompt.md", "PROMPT.md", "problem_statement.md", "original_prompt.txt"):
        if (d / f).exists():
            text = read_narrative(d / f)
            break
    return {"forums": classify(text, FORUM_RULES), "events": classify(text, EVENT_RULES)}


def org_family(card):
    if card.get("org_family"):
        return card["org_family"]
    hits = classify(((card.get("world") or {}).get("org_type") or ""), ORG_RULES)
    return hits[0] if hits else None
