"""The contractor's and the Authority's records as they hold them: the survey's two deliveries (an autumn-built
workbook for the spring and summer fieldwork, a CSV for the autumn and winter), the expansion weights, the codebook
and the Authority's settlement gazetteer."""
from __future__ import annotations

import csv
import datetime as dt

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import common as K
import world as W

INTERVIEWERS = ["INT02", "INT03", "INT05", "INT07", "INT08", "INT11", "INT12"]
AGE1 = ["16-24", "25-44", "45-64", "65-74", "75+"]                # wave 1 bands
AGE2 = ["16-34", "35-54", "55-74", "75+"]                         # wave 2 bands
MODES = ["Car or van", "On foot", "Bicycle", "Bus", "Motorcycle", "Other"]
ACTIVITIES = ["Walking", "Dog walking", "Running", "Cycling", "Family outing", "Watersports", "Visiting the hall",
              "Picnic", "Wildlife watching", "Cafe", "Event", "Fishing"]


def _age_years(typ, r):
    if typ == "reg":
        return int(np.clip(r.normal(56, 13), 16, 92))
    if typ == "loc":
        return int(np.clip(r.normal(46, 16), 16, 92))
    return int(np.clip(r.normal(43, 15), 16, 90))


def _band(age, bands):
    for b in bands:
        lo, hi = (int(b.rstrip("+")), 200) if b.endswith("+") else map(int, b.split("-"))
        if lo <= age <= hi:
            return b
    raise ValueError(age)


def interviews(T):
    """One record per completed interview, in fieldwork order: who was interviewed (their home, kind and
    reported 12-month visit count) and the survey's own fields."""
    P, S = T["P"], T["S"]
    rr = W.rng("report", 35)        # the reporting stream, drawn so no estimate lands on a round hundred
    ra = W.rng("attrs")
    out = []
    for j, (pk, s, d) in enumerate(K.STRATA):
        dd, md = S["days"][j], S["m_day"][j]
        picks = list(S["picks"][j])
        k = 0
        for day, m in zip(dd, md):
            iv = sorted(ra.choice(INTERVIEWERS, 1)) * m
            starts = np.sort(ra.integers(9 * 60 + 35, 16 * 60 + 40, m))
            for t in range(m):
                idx = int(picks[k])
                k += 1
                code, typ, f, nz, cnt = P[idx]
                per = {}
                for jj, c in zip(nz, cnt):
                    per[K.STRATA[int(jj)][0]] = per.get(K.STRATA[int(jj)][0], 0) + int(c)
                # Q9 asks for each park in turn: each answer as people give it, the park they are leaving at least 1
                vp = {q: (W.report(per[q], rr) if per.get(q) else 0) for q in K.PARK_CODES}
                vp[pk] = max(vp[pk], 1)
                age = _age_years(typ, ra)
                place = (K.OUTSIDE_PLACES[int(ra.integers(len(K.OUTSIDE_PLACES)))] if code == "OUT" else
                         K.SETTLEMENTS[code][int(min(ra.geometric(.45) - 1, 5))])
                dog = ra.random() < (.72 if typ == "reg" else .18 if typ == "loc" else .09)
                car = ra.random() < (.38 if typ == "reg" else .62 if typ == "loc" else .9)
                mode = "Car or van" if car else str(ra.choice(MODES[1:], p=[.62, .2, .1, .03, .05]))
                if dog:
                    act = "Dog walking"
                elif pk == "DW" and ra.random() < .45:
                    act = "Watersports"
                elif pk == "TH" and ra.random() < .5:
                    act = "Visiting the hall"
                else:
                    act = str(ra.choice(ACTIVITIES, p=[.3, 0, .07, .07, .2, .03, .02, .08, .07, .1, .03, .03]))
                adults = int(1 if typ == "reg" and ra.random() < .7 else min(1 + ra.poisson(.9), 6))
                kids = int(0 if typ == "reg" else ra.poisson(.55 if typ in ("occ", "trip") else .4))
                out.append(dict(park=pk, season=s, daytype=d, date=day, start=int(starts[t]), interviewer=iv[t],
                                code=code, kind=typ, f_true=f, f_here=per[pk], vp=vp, visits=sum(vp.values()),
                                age=age, place=place,
                                sex=str(ra.choice(["Female", "Male", "Prefer not to say"], p=[.53, .46, .01])),
                                adults=adults, children=kids, dog=dog, mode=mode, activity=act,
                                rating=int(ra.choice([2, 3, 4, 5], p=[.03, .12, .41, .44])), stratum=j, person=idx))
    # some respondents give no usable home (refused, or "between homes"); the contractor codes them NK
    rn = W.rng("nk")
    for o in out:
        if rn.random() < .006:
            o["code"], o["place"] = "NK", str(rn.choice(["Refused", "No fixed address", "Travelling", ""]))
    out.sort(key=lambda o: (o["date"], o["park"], o["start"]))
    n1 = n2 = 0
    for o in out:
        if o["season"] in ("SP", "SU"):
            n1 += 1
            o["wave"], o["serial"] = 1, n1
        else:
            n2 += 1
            o["wave"], o["serial"] = 2, n2
    return out


def weights_table(T, recs):
    S = T["S"]
    m = np.zeros(len(K.STRATA), int)
    for o in recs:
        m[o["stratum"]] += 1
    est = S["est"]
    for o in recs:
        j = o["stratum"]
        o["weight"] = round(float(est[j]) / int(m[j]), 6)
        o["stratum_est"], o["stratum_n"] = int(est[j]), int(m[j])
    return m


def _xl(ws, header, rows, widths, fmt=None):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5D50")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i) if i < 26 else "A" + chr(65 + i - 26)].width = w
    ws.freeze_panes = "A2"
    if fmt:
        for col, f in fmt.items():
            for cell in ws[col][1:]:
                cell.number_format = f


def write_wave1(path, recs):
    wb = Workbook()
    ws = wb.active
    ws.title = "Interviews"
    header = ["Interview No", "Park", "Interview date", "Start", "Interviewer", "Home town/village", "Home area",
              "Age group", "Sex", "Adults in party", "Children in party", "Dog with party", "Arrived by",
              "Main activity"] + ["Visits, last 12 months (incl. today): %s" % K.PARK_NAME[q] for q in K.PARK_CODES] + [
              "Overall rating (1-5)"]
    rows = []
    for o in recs:
        if o["wave"] != 1:
            continue
        rows.append(["DV25-1-%04d" % o["serial"], K.PARK_NAME[o["park"]], o["date"],
                     "%02d:%02d" % divmod(o["start"], 60), o["interviewer"], o["place"] or None, o["code"],
                     _band(o["age"], AGE1), o["sex"], o["adults"], o["children"], "Yes" if o["dog"] else "No",
                     o["mode"], o["activity"]] + [o["vp"][q] for q in K.PARK_CODES] + [o["rating"]])
    _xl(ws, header, rows, [14, 24, 13, 7, 11, 20, 10, 10, 16, 9, 9, 9, 13, 18] + [13] * 6 + [9], {"C": "dd/mm/yyyy"})
    ws2 = wb.create_sheet("Notes")
    for line in ["Daskbeck Valley Regional Park Authority visitor survey 2025-26, wave 1 (fieldwork 1 April to 30 "
                 "September 2025).", "Delivered by %s, 14 October 2025. One row per completed interview." % K.CONTRACTOR,
                 "Home area: the Authority's area code for the member council the respondent lives in, from the "
                 "settlement gazetteer; OUT lives outside the nine member councils; NK not known.",
                 "Weights are supplied separately (survey_weights_2025-26.csv)."]:
        ws2.append([line])
    ws2.column_dimensions["A"].width = 110
    wb.save(path)
    return len(rows)


def write_wave2(path, recs):
    header = ["case_ref", "site_code", "int_date", "int_time", "int_id", "home_place", "home_area_code", "age_band",
              "gender", "n_adults", "n_children", "dog", "mode", "activity"] + ["q9_%s" % q.lower() for q in K.PARK_CODES] + [
              "q14_rating"]
    n = 0
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        for o in recs:
            if o["wave"] != 2:
                continue
            n += 1
            w.writerow(["DV25/2/%04d" % o["serial"], o["park"], o["date"].strftime("%d/%m/%Y"),
                        "%02d:%02d" % divmod(o["start"], 60), o["interviewer"], o["place"], o["code"],
                        _band(o["age"], AGE2), {"Female": "F", "Male": "M"}.get(o["sex"], "X"), o["adults"],
                        o["children"], "Y" if o["dog"] else "N", o["mode"].split(" ")[0].upper(),
                        o["activity"].upper().replace(" ", "_")] + [o["vp"][q] for q in K.PARK_CODES] + [o["rating"]])
    return n


def write_weights(path, recs):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["case_id", "stratum", "adult_visits_est", "interviews_in_stratum", "weight"])
        for o in sorted(recs, key=lambda o: (o["wave"], o["serial"])):
            pk, s, d = K.STRATA[o["stratum"]]
            w.writerow(["%d-%04d" % (o["wave"], o["serial"]), "%s/%s/%s" % (pk, s, d), o["stratum_est"],
                        o["stratum_n"], "%.6f" % o["weight"]])
    return len(recs)
