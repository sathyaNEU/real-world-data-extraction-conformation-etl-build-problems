"""The pack's documents: the GPI methodology statement (2026 edition), the index committee's terms of reference for the
2027 structure, the partner register, the feed dictionary, the household panel's format split by category, and the
export manifest."""
from __future__ import annotations

import csv
import datetime as dt

from reportlab import rl_config

rl_config.invariant = 1
from openpyxl import Workbook                                                       # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill                            # noqa: E402
from reportlab.lib import colors                                                    # noqa: E402
from reportlab.lib.pagesizes import A4                                              # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                     # noqa: E402
from reportlab.lib.units import mm                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                            # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                        # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import common as K   # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("Sans", "LiberationSans-Regular"), ("Sans-B", "LiberationSans-Bold"), ("Sans-I", "LiberationSans-Italic")):
    pdfmetrics.registerFont(TTFont(nm, FONT_DIR + f + ".ttf"))
pdfmetrics.registerFontFamily("Sans", normal="Sans", bold="Sans-B", italic="Sans-I", boldItalic="Sans-B")
INK, GREY, RULE = colors.HexColor("#1f2428"), colors.HexColor("#5f6870"), colors.HexColor("#b9c0c7")
ACCENT = colors.HexColor("#7a2e3b")


def styles():
    font = "Sans"
    body = ParagraphStyle("b", fontName=font, fontSize=9.6, leading=13.6, textColor=INK, spaceAfter=5)
    return dict(
        body=body,
        h1=ParagraphStyle("h1", fontName=font + "-B", fontSize=14.5, leading=18.5, textColor=ACCENT, spaceAfter=2),
        h2=ParagraphStyle("h2", fontName=font + "-B", fontSize=10.4, leading=14, textColor=INK, spaceBefore=7,
                          spaceAfter=2),
        meta=ParagraphStyle("m", fontName=font, fontSize=8.3, leading=11.2, textColor=GREY, spaceAfter=9),
        note=ParagraphStyle("n", fontName=font, fontSize=7.9, leading=10.4, textColor=GREY, spaceBefore=6),
        cell=ParagraphStyle("c", fontName=font, fontSize=8.6, leading=11, textColor=INK),
        cellb=ParagraphStyle("cb", fontName=font + "-B", fontSize=8.6, leading=11, textColor=INK))


def _doc(path, title):
    return SimpleDocTemplate(path, pagesize=A4, leftMargin=21 * mm, rightMargin=21 * mm, topMargin=21 * mm,
                             bottomMargin=18 * mm, title=title, author=K.FIRM, subject="", creator="", producer="")


def _frame(left):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(GREY)
        canvas.setFont("Sans", 7.4)
        canvas.drawString(21 * mm, A4[1] - 11.5 * mm, left)
        canvas.drawRightString(A4[0] - 21 * mm, A4[1] - 11.5 * mm, K.FIRM)
        canvas.drawRightString(A4[0] - 21 * mm, 10 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _table(S, rows, widths, right=()):
    data = []
    for i, r in enumerate(rows):
        out = []
        for j, c in enumerate(r):
            st = S["cellb"] if i == 0 else S["cell"]
            if j in right and i > 0:
                st = ParagraphStyle("r", parent=st, alignment=2)
            out.append(Paragraph(str(c), st))
        data.append(out)
    t = Table(data, colWidths=[w * mm for w in widths], hAlign="LEFT", repeatRows=1)
    t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, 0), 0.7, INK), ("LINEBELOW", (0, 1), (-1, -1), 0.25, RULE),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 2.2),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2)]))
    return t


# ---------------------------------------------------------------------------------------------- methodology
def write_method(path):
    S = styles()
    doc = _doc(path, "Delvett Grocery Price Index: methodology statement, 2026 edition")
    st = [Paragraph("Delvett Grocery Price Index: methodology statement", S["h1"]),
          Paragraph("2026 edition, for retail year 2026 (the weeks ending 6 December 2025 to 28 November 2026). "
                    "Price Indices, %s. Issued 28 January 2026 by %s, Head of Price Indices; replaces the 2025 edition."
                    % (K.FIRM, K.PEOPLE["head"]), S["meta"])]
    parts = [
        ("1. What the index measures",
         "The Delvett Grocery Price Index (GPI) measures the change in the prices shoppers pay for the same grocery "
         "items at the panel's stores. It moves only when prices move: a change in how much shoppers buy, in which "
         "items they buy or in which stores they buy them is not a price change and does not by itself move the "
         "index. Each item counts in proportion to what shoppers spend on it."),
        ("2. The panel",
         "The panel is the stores of the firm's partner chains in four store formats: supermarkets, neighbourhood "
         "stores, convenience stores and online, whose stores are the chains' fulfilment sites. A chain's stores count "
         "in the panel from the first week its data agreement covers, as the partner register gives it, and each "
         "store from the week the panel store register gives for it."),
        ("3. The basket",
         "The 2026 basket is the lines in the item master: sixteen lines in each of eighteen food categories, chosen "
         "at the 2025 basket review and held for the year. Lines sold by weight are priced by the kilogram."),
        ("4. Prices and spending",
         "The chains send movement data each week: for each store, line and week, the units sold (kilograms for lines "
         "sold by weight) and the dollars taken at the till after any discount. Prices are dollars taken over units "
         "sold, and spending is dollars taken. The loader places each chain's trading week on the firm's retail week, "
         "Sunday to Saturday."),
        ("5. Comparing weeks within a format",
         "Within each store format the index for a category compares every pair of weeks in the retail year. The "
         "comparison of two weeks is a Tornqvist index over the products sold in both weeks: the geometric mean of "
         "their price relatives, each weighted by the mean of its shares of the spending on those products in the two "
         "weeks. The GEKS method combines the comparisons: the index for week t is the geometric mean, over every "
         "week s of the year, of the comparison of the year's first week with week s multiplied by the comparison of "
         "week s with week t, so that the first week's index is 1."),
        ("6. The change over the year",
         "A category's change over the year in a format is the average of its index over the year's last four weeks "
         "divided by its average over the year's first four weeks, less one, in per cent."),
        ("7. National figures",
         "A category's national change is the weighted average of its four format changes, with the format weights "
         "the index committee sets for the year."),
        ("8. Data as loaded",
         "The movement table keeps every load a chain sends. For each store and week the index uses the chain's "
         "accepted load, and a resubmission replaces the loads it covers; the feed delivery log records each load's "
         "status. Item codes are matched to the item master as GTIN-14."),
        ("9. Publication",
         "The GPI is published weekly, nationally and by format, for all food and for the categories the committee "
         "publishes on their own. Published figures are not revised after four weeks."),
    ]
    for h, t in parts:
        st.append(Paragraph(h, S["h2"]))
        st.append(Paragraph(t, S["body"]))
    st.append(Paragraph("Enquiries: Price Indices, %s." % K.FIRM, S["note"]))
    doc.build(st, onFirstPage=_frame("GPI methodology statement, 2026 edition"),
              onLaterPages=_frame("GPI methodology statement, 2026 edition"))


# ---------------------------------------------------------------------------------------------- the committee's terms
def write_terms(path):
    S = styles()
    doc = _doc(path, "Index Committee: terms of reference for the 2027 GPI publication structure")
    st = [Paragraph("Index Committee: terms of reference for the 2027 publication structure", S["h1"]),
          Paragraph("Agreed by the Index Committee on 4 November 2026. The committee adopts the 2027 structure at its "
                    "meeting on Wednesday 9 December 2026.", S["meta"])]
    parts = [
        ("1. Purpose",
         "In 2027 the GPI publishes a national food index and, for each food category whose change holds across the "
         "store formats, a stand-alone sub-index. The other categories are published together as Other food."),
        ("2. The test",
         "A food category is published as a stand-alone sub-index in 2027 if, over the 52 retail weeks ending Saturday "
         "6 December 2025 to Saturday 28 November 2026, its national change is within 1.5 percentage points of its "
         "change in each of the four store formats."),
        ("3. Measurement",
         "Changes are measured as the GPI methodology statement (2026 edition) sets out, from the movement data the "
         "partner chains had delivered by Friday 4 December 2026. The national change combines the four format changes "
         "with the committee's 2026 format weights in Table 1."),
        ("4. Categories",
         "The eighteen food categories of the 2026 basket, as the item master gives them."),
        ("5. Papers",
         "The price index team circulates the proposed structure, its supporting figures and a paper for the "
         "committee by Monday 7 December 2026."),
    ]
    for h, t in parts:
        st.append(Paragraph(h, S["h2"]))
        st.append(Paragraph(t, S["body"]))
    st.append(Spacer(1, 5))
    st.append(Paragraph("Table 1. The committee's format weights for 2026", S["cellb"]))
    st.append(Spacer(1, 2))
    rows = [["Store format", "Weight (per cent)"]] + [[K.FORMAT_LABEL[f], "%.1f" % K.FORMAT_WEIGHTS[f]]
                                                      for f in K.FORMATS] + [["Total", "100.0"]]
    st.append(_table(S, rows, [60, 34], right=(1,)))
    st.append(Paragraph("Shares of household grocery spending by store format in retail year 2025, from the firm's "
                        "household panel, set at the committee's weights review of 12 February 2026.", S["note"]))
    doc.build(st, onFirstPage=_frame("Index Committee: terms of reference, 2027 structure"),
              onLaterPages=_frame("Index Committee: terms of reference, 2027 structure"))


# ---------------------------------------------------------------------------------------------- workbooks
HEAD = Font(name="Calibri", bold=True, size=11)
FILL = PatternFill("solid", fgColor="DDE3EA")


def _sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]:
        c.font = HEAD
        c.fill = FILL
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"


PRICE_CODES = {
    "S1": [("SPECIAL", "Catalogue special (temporary price reduction)")],
    "S2": [("PROMO", "Catalogue promotion, reduced shelf price for the week")],
    "S3": [("TPR", "Temporary price reduction")],
    "S4": [("SALE", "Weekly sale price")],
    "S5": [("OFFER", "Special offer price for the week")],
    "N1": [("P", "Promotional price (catalogue)"), ("RC", "Reduced to clear (store markdown, short-dated stock)")],
    "N2": [("PROMO", "Catalogue promotion"), ("MD", "Store markdown, short-dated or damaged stock")],
    "N3": [("SPC", "Special (temporary price reduction)"), ("RC", "Reduced to clear")],
    "C1": [("PR", "Promotional price for the week"), ("RC", "Reduced to clear")],
    "C2": [("TPR", "Temporary price reduction (catalogue)"), ("PC", "Price check override at the register"),
           ("RC", "Reduced to clear")],
    "C3": [("DEAL", "Weekly deal price"), ("CLR", "Clearance markdown by the store")],
    "O1": [("SALE", "Sale price for the week")],
    "O2": [("PROMO", "Online promotion price")],
    "O3": [("SPECIAL", "Weekly special")],
}
CODE_LIST_DATE = {"S1": "2025-03-11", "S2": "2024-11-20", "S3": "2025-07-02", "S4": "2024-09-16", "S5": "2026-06-30",
                  "N1": "2025-05-21", "N2": "2024-12-09", "N3": "2025-10-14", "C1": "2025-02-26", "C2": "2026-08-03",
                  "C3": "2024-10-28", "O1": "2025-08-19", "O2": "2025-06-05", "O3": "2026-01-22"}


def write_partners(path, agreement):
    wb = Workbook()
    ws = wb.active
    ws.title = "Chains"
    rows = []
    for c in K.CHAIN_CODES:
        ch = K.CH[c]
        rows.append([c, ch["banner"], K.FORMAT_LABEL[ch["fmt"]], agreement[c].strftime("%Y-%m-%d"),
                     "Monday to Sunday" if c == K.SUNDAY_CHAIN else "Sunday to Saturday",
                     "One shelf price for each line in all the chain's %s stores, set by head office and changed "
                     "weekly; catalogue specials run in every store for the week" % K.FORMAT_LABEL[ch["fmt"]].lower(),
                     CODE_LIST_DATE[c]])
    _sheet(ws, ["chain_code", "banner", "format", "data_agreement_from", "trading_week", "shelf_pricing",
                "code_list_supplied"], rows, [11, 22, 15, 19, 19, 92, 18])
    for r in ws.iter_rows(min_row=2):
        r[5].alignment = Alignment(wrap_text=True, vertical="top")
    ws2 = wb.create_sheet("Price codes")
    rows = [[c, code, meaning] for c in K.CHAIN_CODES for code, meaning in PRICE_CODES[c]]
    _sheet(ws2, ["chain_code", "code", "meaning"], rows, [11, 10, 58])
    wb.save(path)


DICTIONARY = [
    ("gpi_movement_ry2026.csv", "", "", "One row for each store, line and week in which the store sold the line, for each "
     "load received. A store that sold none of a line in a week has no row for it that week."),
    ("gpi_movement_ry2026.csv", "load_id", "text", "The delivery the row came in; the feed delivery log gives each "
     "load's status. The table keeps every load received."),
    ("gpi_movement_ry2026.csv", "store_id", "text", "The firm's store code (panel store register)."),
    ("gpi_movement_ry2026.csv", "item_code", "text", "The line's GTIN as the chain's feed sends it."),
    ("gpi_movement_ry2026.csv", "week_ending", "date", "Saturday ending the firm's retail week. The loader places a "
     "chain's trading week on the retail week holding six of its seven days."),
    ("gpi_movement_ry2026.csv", "units", "number", "Units sold; kilograms (three decimals) for lines whose unit of "
     "sale is kg."),
    ("gpi_movement_ry2026.csv", "sales_aud", "number", "Dollars taken at the till for the row's units, after any "
     "discount, including GST where it applies."),
    ("gpi_movement_ry2026.csv", "promo_code", "text", "The till's price code on the row's sales, in the chain's own "
     "code list (partner register, Price codes); blank for sales at the shelf price."),
    ("store_weekly_sales_ry2026.csv", "", "", "Each store's weekly takings by department, for each load received."),
    ("store_weekly_sales_ry2026.csv", "load_id", "text", "The delivery the row came in (feed delivery log)."),
    ("store_weekly_sales_ry2026.csv", "store_id", "text", "The firm's store code."),
    ("store_weekly_sales_ry2026.csv", "week_ending", "date", "The last day of the chain's trading week, as the chain "
     "sends it."),
    ("store_weekly_sales_ry2026.csv", "dept_code", "text", "The chain's department code."),
    ("store_weekly_sales_ry2026.csv", "dept_name", "text", "The chain's department name."),
    ("store_weekly_sales_ry2026.csv", "sales_aud", "number", "Sales in dollars, including GST, as the chain reports "
     "them."),
    ("panel_store_register.csv", "", "", "One row for each store code the firm has issued to a panel store."),
    ("panel_store_register.csv", "store_id", "text", "The firm's store code."),
    ("panel_store_register.csv", "chain_code", "text", "The partner chain (partner register)."),
    ("panel_store_register.csv", "format", "text", "The store's format."),
    ("panel_store_register.csv", "store_name", "text", "The store's trading name as the chain gives it."),
    ("panel_store_register.csv", "suburb, state, postcode", "text", "The store's location."),
    ("panel_store_register.csv", "first_week_ending", "date", "The first week the store's data counts in the panel."),
    ("panel_store_register.csv", "last_week_ending", "date", "The last week the store's data counted in the panel; "
     "blank while it continues."),
    ("gpi_item_master_2026.csv", "", "", "One row for each line in the 2026 basket."),
    ("gpi_item_master_2026.csv", "item_code", "text", "GTIN-14."),
    ("gpi_item_master_2026.csv", "category_code, category", "text", "The line's food category."),
    ("gpi_item_master_2026.csv", "unit_of_sale", "text", "each, or kg for lines sold by weight."),
    ("feed_delivery_log_ry2026.csv", "", "", "One row for each load received for retail year 2026."),
    ("feed_delivery_log_ry2026.csv", "first_week_ending, last_week_ending", "date", "The weeks the load covers, in the "
     "chain's own week-ending dates."),
    ("feed_delivery_log_ry2026.csv", "received_utc", "timestamp", "When the load reached the firm (UTC)."),
    ("feed_delivery_log_ry2026.csv", "rows", "integer", "Rows written to the table from the load."),
    ("feed_delivery_log_ry2026.csv", "status", "text", "accepted; superseded (replaced by a later resubmission); "
     "rejected (failed validation; the table keeps its rows)."),
    ("retail_calendar_2025-2027.csv", "", "", "One row for each retail week of retail years 2025 to 2027."),
    ("retail_calendar_2025-2027.csv", "retail_year, retail_week", "integer", "The firm's retail year (weeks ending "
     "on the Saturday nearest the end of November) and its week number."),
    ("retail_calendar_2025-2027.csv", "week_start, week_ending", "date", "Sunday and Saturday of the retail week."),
]
TERMS = [
    ("share of spending on promotion", "Spending on promotion as a share of all spending on the lines in question at "
     "the panel's stores over the period."),
    ("spending on promotion", "Spending at a catalogue special: a temporary price reduction a chain runs in all its "
     "stores in a format for the week. Store markdowns of short-dated stock and price check overrides at the register "
     "are not promotions."),
    ("store format", "Supermarket, neighbourhood, convenience or online, as the panel store register gives it."),
    ("retail week", "Sunday to Saturday (retail calendar)."),
]


def write_dictionary(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Fields"
    _sheet(ws, ["file", "field", "type", "description"], [[v if v else None for v in r] for r in DICTIONARY],
           [32, 34, 11, 104])
    for r in ws.iter_rows(min_row=2):
        r[3].alignment = Alignment(wrap_text=True, vertical="top")
    ws2 = wb.create_sheet("Terms")
    _sheet(ws2, ["term", "meaning"], [list(t) for t in TERMS], [26, 110])
    for r in ws2.iter_rows(min_row=2):
        r[1].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def write_household(path):
    """The household panel's split of each category's spending across the store formats, retail weeks to 31 October
    2026: the firm's consumer panel, a different source from the partner chains' tills."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Format split"
    r = K.rng("household")
    rows = []
    base = [0.60, 0.16, 0.07, 0.17]
    for cat in K.CATS:
        v = [b * float(r.uniform(0.82, 1.18)) for b in base]
        s = sum(v)
        v = [round(100 * x / s, 1) for x in v]
        v[0] = round(100 - sum(v[1:]), 1)
        rows.append([K.CAT_CODE[cat], cat] + v + [int(r.integers(5200, 8300))])
    _sheet(ws, ["category_code", "category", "supermarket_pct", "neighbourhood_pct", "convenience_pct", "online_pct",
                "buying_households"], rows, [15, 24, 17, 19, 17, 12, 19])
    ws2 = wb.create_sheet("About")
    ws2.append(["Household panel: share of each food category's spending by store format"])
    ws2.append(["Retail weeks ending 6 December 2025 to 31 October 2026. 8,400 panel households scanning their "
                "grocery purchases at home; shares weighted to the household population."])
    ws2.append(["Consumer Panels team, %s. Extracted 10 November 2026." % K.FIRM])
    ws2.column_dimensions["A"].width = 120
    wb.save(path)


def write_manifest(path, files, counts):
    desc = {
        "movement": ("Movement table, retail year 2026, every load", "Index loader (scanner movement)"),
        "totals": ("Store takings by department, retail year 2026, every load", "Store sales loader"),
        "register": ("Panel store register", "Partner data team"),
        "partners": ("Partner register: chains and price code lists", "Partner data team"),
        "items": ("2026 basket item master", "Price Indices"),
        "log": ("Feed delivery log, retail year 2026", "Partner data team"),
        "calendar": ("Retail calendar, retail years 2025 to 2027", "Price Indices"),
        "method": ("GPI methodology statement, 2026 edition", "Price Indices"),
        "terms": ("Index Committee terms of reference, 2027 structure", "Index Committee secretariat"),
        "dictionary": ("Feed dictionary", "Partner data team"),
        "household": ("Household panel format split by category, retail year 2026 to October", "Consumer Panels"),
        "audit": ("Shelf label price checks, 2025", "Field audit team"),
    }
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(["file", "contents", "rows_or_pages", "source", "extracted"])
        for k, name in files.items():
            if k == "manifest":
                continue
            d, src = desc[k]
            w.writerow([name, d, counts.get(k, ""), "%s, %s" % (src, K.FIRM), K.EXPORT.strftime("%Y-%m-%d")])
        w.writerow(["", "Folder compiled by Price Indices, %s, %s. Internal." % (
            K.FIRM, K.EXPORT.strftime("%-d %B %Y")), "", "", ""])
