"""The build's assertions, on the shipped pack and on the generator's clean truth: the three-rules block (R1 to R3c),
the ladder and its separation, the semantic fork table (SF1 to SF13), the closure moves, the asks' devices and stops,
the cracker sheet, the main population's separation from every device, the distractors' deletion test, the input
gates, the realism sweep and the H21 name screen."""
from __future__ import annotations

import collections
import datetime as dt
import json
import os
import re
import shutil
import tempfile
import zipfile

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from pypdf import PdfReader

import common as K
import deliver as DL
import golden as G
import ops as O
import world as W

PASSED = []


def check(label, ok):
    assert ok, "check failed: " + label
    PASSED.append(label)


def share(N, codes=K.VALLEY):
    return 100 * sum(N[c] for c in codes) / sum(N[c] for c in K.CODES)


def levies(N):
    return G.rounded(G.levy(N))


# ---------------------------------------------------------------------------------------------- shipped text
def pdf_text(path):
    return " ".join((p.extract_text() or "") for p in PdfReader(path).pages)


def xlsx_text(path):
    wb = load_workbook(path, read_only=True)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            out.extend(str(v) for v in row if isinstance(v, str))
    return " ".join(out)


def texts(target):
    out = {}
    for f in sorted(os.listdir(target)):
        p = os.path.join(target, f)
        if f.endswith(".pdf"):
            out[f] = pdf_text(p)
        elif f.endswith(".xlsx"):
            out[f] = xlsx_text(p)
        else:
            out[f] = open(p, encoding="utf-8").read()
    return out


def flat(s):
    return re.sub(r"\s+", " ", s)


# ---------------------------------------------------------------------------------------------- the clean truth
def clean(T):
    """The world before any record is written: who each interview was (home, true visits), and each stratum's true
    adult visits over its interviews as the weight."""
    m = T["m"]
    rows = [(o["code"] if o["code"] != "NK" else T["P"][o["person"]][0], o["f_true"], T["S"]["V"][o["stratum"]] / m[o["stratum"]],
             o["f_here"]) for o in T["recs"]]
    est = {}
    for how in ("raw", "visits", "mean", "unweighted", "golden", "here"):
        s = collections.defaultdict(float)
        wf = collections.defaultdict(float)
        for code, f, w, fh in rows:
            if code not in K.CODES:
                continue
            s[code] += {"raw": 1, "visits": w, "unweighted": 1 / f, "golden": w / f, "mean": w, "here": w / fh}[how]
            wf[code] += w * f
        if how == "mean":
            s = {c: s[c] / (wf[c] / s[c]) for c in K.CODES}
        est[how] = {c: s[c] for c in K.CODES}
    truth = collections.Counter(p[0] for p in T["P"])
    est["truth"] = {c: float(truth[c]) for c in K.CODES}
    return est


# ---------------------------------------------------------------------------------------------- helpers on a copy
def copy_without(target, drop):
    tmp = tempfile.mkdtemp()
    for f in os.listdir(target):
        if f not in drop:
            shutil.copy2(os.path.join(target, f), tmp)
    return tmp


def all_figures(target):
    o = G.compute(target)
    return (G.rounded(o["levy"]), {p: round(v, 2) for p, v in o["income"].items()},
            {c: (v["on_time"], v["late"]) for c, v in o["collection"].items()}, o["vehicles"])


def miss(a, b, tol=0.5):
    return abs(a - b) > tol


def lev(a, b):
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


# ---------------------------------------------------------------------------------------------- run
def run(target, T):
    F = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "figures.json")))
    o = G.compute(target)
    d = o["survey"]
    N = o["N"]
    L = o["levy_r"]
    tx = texts(target)
    papers = flat(tx[K.FILES["levypapers"]])
    note = flat(tx[K.FILES["fieldnote"]])
    E = clean(T)

    # ------------------------------------------------------------------------------------------ the three rules
    wt = pd.read_csv(G.f(target, "weights"))
    check("R1 the field's raw material ships: one row per interview with its home area, its 12-month visits to each "
          "of the six parks and its expansion weight (%d interviews)" % len(d), len(d) == len(T["recs"]) and
          d[["v_" + q for q in K.PARK_CODES]].notna().all().all() and (d["here"] >= 1).all())
    check("R1 no shipped column holds the deciding quantity: the weights file carries visits only (weight = stratum "
          "adult visits over its interviews, to six places)", list(wt.columns) == ["case_id", "stratum", "adult_visits_est",
                                                                                    "interviews_in_stratum", "weight"]
          and (abs(wt["weight"] - wt["adult_visits_est"] / wt["interviews_in_stratum"]) < 6e-7).all())
    check("R1 no shipped file states a resident count, a share or a split on the new basis",
          not any(re.search(r"residents? who visited", t, re.I) and re.search(r"\d{2},\d{3} residents", t)
                  for t in tx.values()))
    check("R1 the golden's decisive computation reads each interview's weight and its own visit count",
          abs(sum(N.values()) - float((d.loc[d["code"].isin(K.CODES), "weight"] /
                                       d.loc[d["code"].isin(K.CODES), "visits"]).sum())) < 1e-6)
    # R2: each Layer 1 item, naturally mishandled, leaves the call (or fails loudly)
    w2 = pd.read_csv(G.f(target, "wave2"), dtype=str)
    md = pd.to_datetime(w2["int_date"], format="%d/%m/%Y")
    check("R2 dates in two formats: the estimator does not read the interview date, so a month-first misreading of "
          "wave 2's dates leaves every levy where it was", levies(G.residents(d.assign(date=None))) == L and len(md) == len(w2))
    serial_only = d[["case"]].assign(k=d["case"].str[2:]).merge(wt.assign(k=wt["case_id"].str[2:]), on="k")
    check("R2 interview numbers written three ways: joining the weights on the serial alone fans out (%d rows from %d), "
          "a loud failure, never a silent one" % (len(serial_only), len(d)), len(serial_only) > len(d) * 1.4)
    check("R2 home areas coded NK: dropping them or sharing them out pro rata gives the same levies",
          levies({c: N[c] * (1 + 0.0) for c in K.CODES}) == L and
          levies({c: N[c] + float((d["code"] == "NK").sum()) * N[c] / sum(N.values()) for c in K.CODES}) == L)
    heap = E["golden"]
    check("R2 heaped visit counts: the estimator on the answers as given and on the true counts gives the valley "
          "councils' share within 2 points (%.2f against %.2f)" % (share(N), share(heap)), abs(share(N) - share(heap)) < 2)
    check("R2 age bands recoded and parks as names or codes: neither enters the estimator, so dropping both leaves "
          "every levy where it was", levies(G.residents(d.drop(columns=["park"]))) == L)
    # R3a: the analyst test on the clean truth
    check("R3a the default path on the clean truth (true visit counts, true stratum visits) still files a wrong split: "
          "valley share %.1f against the truth's %.1f" % (share(E["mean"]), share(E["truth"])),
          abs(share(E["mean"]) - share(E["truth"])) > 18)
    check("R3a executing every filed rule (the resolution, the note's weights) lands the default path's split on the "
          "shipped pack: %.1f per cent" % o["valley"]["mean"], round(o["valley"]["mean"], 1) == 25.0)
    check("R3a the frequency-aware route that counts only the visits to the park of interview lands a wrong split on "
          "the clean truth too (%.1f against the truth's %.1f): a resident who visits two parks can be met at either"
          % (share(E["here"]), share(E["truth"])), abs(share(E["here"]) - share(E["truth"])) > 3)
    check("R3b the decisive step alone (weight over own visit count) moves the default path to the golden, which sits "
          "within 2 points of the truth on the clean world (%.1f against %.1f)" % (share(E["golden"]), share(E["truth"])),
          abs(share(E["golden"]) - share(E["truth"])) < 2)
    check("R3b each lower rung's single correction leaves a wrong split: raw to weighted (R1) %.1f, raw to inverse "
          "visits unweighted (R3) %.1f, visits to visits over the average (R2) %.1f, weight over the visits to the park "
          "of interview (R4) %.1f; golden %.1f" % (o["valley"]["visits"], o["valley"]["unweighted"], o["valley"]["mean"],
                                                   o["valley"]["here"], o["valley"]["golden"]),
          min(abs(o["valley"][k] - o["valley"]["golden"]) for k in ("visits", "unweighted", "mean", "here")) > 2.0)
    check("R3c no plant: no memo, draft, briefing or stakeholder figure argues a split (no document but the levy "
          "papers mentions the levy's apportionment)",
          [f for f, t in tx.items() if re.search(r"apportion", t, re.I)] == [K.FILES["levypapers"]])
    check("R3c no copied block in the survey: no run of five identical consecutive interviews in either wave",
          not any((d[["code", "visits", "weight"]].iloc[i:i + 5].nunique() == 1).all() for i in range(0, len(d) - 5, 5)))

    # ------------------------------------------------------------------------------------------ ladder
    v = o["valley"]
    check("ladder: R0 %.1f, R1 %.1f, R2 %.1f, R3 %.1f, R4 %.1f, golden %.1f (valley share)" % (
        v["raw"], v["visits"], v["mean"], v["unweighted"], v["here"], v["golden"]),
        (round(v["raw"], 1), round(v["visits"], 1), round(v["mean"], 1), round(v["unweighted"], 1), round(v["here"], 1),
         round(v["golden"], 1)) == (86.3, 85.1, 25.0, 53.4, 55.5, 50.7))
    R4 = levies(o["rungs"]["here"]["N"])
    gap4 = min(abs(R4[c] - L[c]) for c in K.CODES)
    check("separation: the frequency-aware rung that divides by the visits to the park of interview (R4) is %d pounds "
          "or more from the golden for every council" % gap4, gap4 >= 20000)
    R3 = levies(o["rungs"]["unweighted"]["N"])
    R2 = levies(o["rungs"]["mean"]["N"])
    gap3 = min(abs(R3[c] - L[c]) for c in K.CODES)
    check("separation: the nearest wrong rung (R3) is %d pounds or more from the golden for every council" % gap3,
          gap3 >= 7000)
    gap2 = min(abs(R2[c] - L[c]) for c in K.VALLEY)
    check("separation: the default path is %d pounds or more from the golden for each valley council" % gap2,
          gap2 >= 300000)
    check("dominance: the default path's error in valley share is eight times the nearest wrong rung's or more",
          abs(v["mean"] - v["golden"]) >= 8 * abs(v["unweighted"] - v["golden"]))
    check("no rung names the golden split: every rung differs from the golden for every council by 1,000 pounds or more",
          all(min(abs(levies(o["rungs"][k]["N"])[c] - L[c]) for c in K.CODES) >= 1000
              for k in ("raw", "visits", "mean", "unweighted", "here")))
    check("the levies add to the requirement: %d" % sum(L.values()), sum(L.values()) == K.LEVY_2027)
    check("no two councils' estimates tie (axis 11)", len({round(x, 3) for x in N.values()}) == 9)

    # ------------------------------------------------------------------------------------------ fork table
    check("SF1 member councils: the levy papers' schedule and the codebook name the same nine councils, and every "
          "home area is one of them, OUT or NK", all(K.NAME[c] in papers for c in K.CODES) and
          set(d["code"]) <= set(K.CODES) | {"OUT", "NK"})
    check("SF2 residents aged 16 and over: the resolution says so and the survey interviews adults only",
          "residents aged 16 and over" in papers and "adults aged 16 and over" in note)
    check("SF3 each resident counted once: the resolution's words", "each resident counted once however many times they "
          "visited" in papers)
    check("SF4 survey year: every interview between 1 April 2025 and 31 March 2026 (%s to %s)" % (d["date"].min(),
                                                                                                  d["date"].max()),
          d["date"].min() >= K.YEAR0 and d["date"].max() <= K.YEAR1 and "1 April 2025 to 31 March 2026" in papers)
    check("SF5 one person across the six parks: Q9 asks for the visits to each park and the resolution counts a "
          "resident once however many of the parks they visited", "visited each of these parks" in note and
          "however many of the parks they visited" in papers)
    check("SF6 the delivered weights are the survey's: the note says they are the weights to use",
          "are the weights to use with this survey" in note)
    w1 = pd.read_excel(G.f(target, "wave1"), sheet_name="Interviews")
    check("SF7 Q9 is the survey's only measure of visits (one field per park in each wave, six in all)",
          sum("visit" in c.lower() for c in w1.columns) == 6 and sum(c.startswith("q9_") for c in w2.columns) == 6)
    check("SF7 every respondent gives at least one visit to the park they were leaving, and most of the valley's "
          "regular walkers name a second park (%d of %d interviews name two or more)" % (
              int((d[["v_" + q for q in K.PARK_CODES]].gt(0).sum(axis=1) >= 2).sum()), len(d)),
          (d["here"] >= 1).all() and (d[["v_" + q for q in K.PARK_CODES]].gt(0).sum(axis=1) >= 2).mean() > .3)
    nk = float((d["code"] == "NK").sum())
    check("SF8 home not placed: not counted or shared pro rata, same levies (%d NK interviews)" % nk,
          levies({c: N[c] * (1 + nk / len(d)) for c in K.CODES}) == L)
    with_out = dict(N, OUT=o["outside"])
    t10 = sum(with_out.values())
    check("SF9 visitors from outside: counting them in a tenth share would move every council's levy; paragraph 6 "
          "rules it out", "whose homes lie outside the member councils are not counted" in papers and
          all(abs(K.LEVY_2027 * N[c] / t10 - L[c]) > 1000 for c in K.CODES))
    check("SF10 the levy requirement: the budget table's 7,416,300", "Levy requirement 7,416,300" in papers)
    shares_first = {c: int(np.floor(K.LEVY_2027 * round(N[c] / sum(N.values()), 3) + .5)) for c in K.CODES}
    check("SF11 rounding path: shares rounded first would move the pounds; paragraph 5 rounds each levy",
          shares_first != L and "rounded to the nearest pound" in papers)
    check("SF12 one person interviewed twice: no person identifier; interview numbers unique",
          d["case"].is_unique and not any(re.search(r"person|respondent_id|panel", c, re.I) for c in w1.columns.tolist()
                                         + w2.columns.tolist()))
    party = d.merge(pd.DataFrame({"case": "1-" + w1["Interview No"].str.split("-").str[-1],
                                  "adults": w1["Adults in party"]}).pipe(
        lambda a: pd.concat([a, pd.DataFrame({"case": "2-" + w2["case_ref"].str.split("/").str[-1],
                                              "adults": w2["n_adults"].astype(int)})])), on="case")
    pw = (party["weight"] * party["adults"] / party["visits"]).groupby(party["code"]).sum()
    check("SF13 adults in a group: each answers for themselves (the note); weighting by party size would move the "
          "levies", "answered for themselves only" in note and levies({c: float(pw[c]) for c in K.CODES}) != L)

    # ------------------------------------------------------------------------------------------ closure
    check("axis 9: no interview on 1 April 2025 or 31 March 2026", not d["date"].isin([K.YEAR0, K.YEAR1]).any())
    sh = d.sample(frac=1, random_state=264).reset_index(drop=True)
    check("axes 13 and 14: the estimate is the same over shuffled rows",
          all(abs(G.residents(sh)[c] - N[c]) < 1e-6 for c in K.CODES))
    check("axis 15: interview numbers unique in each wave and in the weights", wt["case_id"].is_unique and
          w1["Interview No"].is_unique and w2["case_ref"].is_unique)
    check("axis 16: every interview joins its weight one to one (%d)" % len(d), len(d) == len(wt) == len(w1) + len(w2))
    check("literal reader: the resolution's estimand is quoted in one file only (the levy papers)",
          sum("whether or not the visitor survey met them" in flat(t) for t in tx.values()) == 1)
    check("literal reader: the every-eighth-adult rule is stated in one file only (the field note)",
          sum("approached every eighth" in flat(t) for t in tx.values()) == 1)
    tr = E["truth"]
    tt = sum(tr.values())
    mae = {}
    for k in ("golden", "raw", "visits", "mean", "unweighted", "here"):
        n = N if k == "golden" else o["rungs"][k]["N"]
        mae[k] = sum(abs(n[c] / sum(n.values()) - tr[c] / tt) for c in K.CODES) / 9 * 100
    check("estimate against the truth: the golden's councils' shares sit closest to the true shares (mean error %.2f "
          "points against %.2f for the nearest rung)" % (mae["golden"], min(v for k, v in mae.items() if k != "golden")),
          mae["golden"] == min(mae.values()))
    check("estimate against the truth: the golden's valley share is the nearest to the truth's (%.1f against %.1f)" % (
        v["golden"], share(tr)), min(("raw", "visits", "mean", "unweighted", "here", "golden"),
                                     key=lambda k: abs(v[k] - share(tr))) == "golden")

    # ------------------------------------------------------------------------------------------ Ask A
    truthA = collections.defaultdict(float)
    for x in T["mrows"]:
        truthA[K.CP_PARK[x[7]]] += x[5] / (1 + K.VAT)
    for ln in T["lines"]:
        truthA[K.CP_PARK[ln["cp"]]] += ln["income"]
    inc = o["income"]
    check("Ask A golden equals the clean record for every park (%s)" % {p: round(inc[p], 2) for p in K.CHARGING},
          all(abs(inc[p] - truthA[p]) < 0.01 for p in K.CHARGING))
    check("Ask A every park's figure sits a quarter pound or more inside its pound (as stated and recomputed from "
          "sessions)", all(abs(inc[p] - np.floor(inc[p]) - .5) >= .25 for p in K.CHARGING))
    app = G.app(target)
    tar = pd.read_excel(G.f(target, "tariffs"), sheet_name="Tariffs 2025-26")
    prices = set(np.round(tar["Charge (GBP)"].astype(float), 2))
    unit = app["Parking income (GBP)"] * 1.2 / app["Sessions"]
    check("Ask A D3: every app line's parking income is sessions times a tariff charge over 1.2 (net of VAT)",
          all(min(abs(u - p) for p in prices) < 0.01 for u in unit))
    check("Ask A D3: the tariffs workbook says charges include VAT; the app statement says nothing of VAT",
          "Charges include VAT at the standard rate" in tx[K.FILES["tariffs"]] and
          "VAT" not in tx[K.FILES["app"]])
    mt = pd.read_csv(G.f(target, "machines"))
    dup = mt.duplicated(["machine_id", "ticket_no"], keep=False)
    check("Ask A D1: three machines' re-sent days (%d rows) keep machine and ticket number under fresh transaction "
          "numbers and a later upload" % int(mt.duplicated(["machine_id", "ticket_no"]).sum()),
          set(mt.loc[dup, "machine_id"]) == {r[0] for r in O.REPLAYS} and mt["txn_id"].is_unique and
          mt.loc[dup].groupby(["machine_id", "ticket_no"])["upload_id"].nunique().eq(2).all())
    check("Ask A D1: the battery's exact-duplicate sweep finds nothing (every row differs in transaction and upload)",
          not mt.duplicated().any())
    reg = pd.read_csv(G.f(target, "assets"))
    check("Ask A D2: the register has PM-2103 at the Water Park, installed 2021, with no move recorded",
          reg.loc[reg["asset_id"] == "PM-2103", "site_code"].item() == "DW-1")
    nat = G.parking_income(target, vat_on_app=True, dedupe=False, sites="register")
    stops = {"natural (VAT off the app income too, every row, register sites)": (nat, 4),
             "VAT off the app income, replays dropped": (G.parking_income(target, vat_on_app=True), 4),
             "app as stated, every row, register sites": (G.parking_income(target, dedupe=False, sites="register"), 3),
             "app as stated, replays dropped, register sites": (G.parking_income(target, sites="register"), 2)}
    for lab, (val, k) in stops.items():
        n = sum(miss(val[p], inc[p], 50) for p in K.CHARGING)
        check("Ask A stop, %s: misses %d of four by 50 pounds or more" % (lab, n), n == k)

    # ------------------------------------------------------------------------------------------ Ask B
    levy_inv = [i for i in T["inv"] if i["kind"] == "levy"]
    sales = {i["no"]: i for i in T["inv"]}
    owner = {i["no"]: i["code"] for i in levy_inv}
    due = {i["no"]: i["due"] for i in levy_inv}
    truthB = {c: dict(on_time=0.0, late=0.0) for c in K.CODES}
    for p in T["pays"]:
        for no, amt in p["covers"]:
            if no in owner and p["value"] <= K.YEAR1:
                truthB[owner[no]]["on_time" if p["value"] <= due[no] else "late"] += amt
    B = o["collection"]
    check("Ask B golden equals the payments as made for every council (on time and late by the day the bank credited "
          "them)", all(abs(B[c][k] - truthB[c][k]) < .005 for c in K.CODES for k in ("on_time", "late")))
    levy25 = {c: sum(i["amount"] for i in levy_inv if i["code"] == c) for c in K.CODES}
    _, i24 = W.levy_2024()
    short = [i["amount"] for i in levy_inv if i["code"] == "CHD" and i["q"] == 3][0] - i24["CHD"][2]
    check("Ask B Chaldale's third instalment at the 2024-25 amount leaves %d pounds unpaid at 31 March" % short,
          short > 0 and abs(levy25["CHD"] - B["CHD"]["on_time"] - B["CHD"]["late"] - short) < .005)
    grb4 = [p for p in T["pays"] if p["payer"] == "GRB" and p["q"] == 4][0]
    check("Ask B Grinbury's fourth instalment reached the bank on %s, after the year and the exports" % grb4["value"],
          grb4["value"] > dt.date(2026, 4, 7))
    cb = pd.read_csv(G.f(target, "cashbook"), dtype=str)
    bk = pd.read_csv(G.f(target, "bank"), dtype=str)
    rows_all = T["rows"]
    inb = [x for x in rows_all if O.EXPORT_FROM <= x["value"] <= O.EXPORT_TO and O.EXPORT_FROM <= x["post"] <= O.CASHBOOK_TO]
    check("Ask B D8: every credit both files hold is posted the working day after the bank credited it (%d credits, one "
          "mechanism, the whole year)" % len(inb), len(inb) > 500 and all(O.next_working(x["value"]) == x["post"]
                                                                          for x in inb))
    key = collections.Counter(zip(bk["Reference"], bk["Amount"]))
    check("Ask B D8: the cash book carries no bank reference; each council's credit is found in the bank's report by "
          "its reference and amount, which no two credits share", "Bank ref" not in cb.columns and
          all(key[(x["ref"], "%.2f" % x["amount"])] == 1 for x in inb if "pay" in x))
    check("Ask B D8: the manifest defines each file's date flatly (the day the finance system booked the receipt, the "
          "day the bank credited the funds) and no file says how one follows the other",
          "the day the finance system booked the receipt" in tx[K.FILES["manifest"]] and
          "the day the bank credited the funds" in tx[K.FILES["manifest"]] and
          not any(re.search(r"morning|next working day|day after|a day later", flat(t), re.I) for t in tx.values()))
    on_due = collections.defaultdict(int)
    for p in T["pays"]:
        for no, _ in p["covers"]:
            if no in owner and p["value"] == due[no]:
                on_due[owner[no]] += 1
    check("Ask B D8: every council has a payment credited on its due date, which the posting date makes late or moves "
          "past the year (%s)" % dict(on_due), all(on_due[c] >= 1 for c in K.CODES))
    council = cb[cb["Invoice"].fillna("") != ""]
    multi = [g for _, g in council.groupby("Receipt no") if len(g) > 1]
    check("Ask B D6: a council's payment run that paid an instalment and a sundry invoice together is one receipt on two "
          "lines with the receipt's amount on each, the two invoices adding to it (%d receipts)" % len(multi),
          len(multi) >= 15 and all(g["Amount"].nunique() == 1 and abs(float(g["Amount"].iloc[0]) - sum(
              sales[n]["amount"] for n in g["Invoice"])) < .005 for g in multi))
    two = collections.defaultdict(set)
    for p in T["pays"]:
        if len(p["covers"]) > 1 and p["post"] <= O.CASHBOOK_TO:
            for no, _ in p["covers"]:
                if no in owner and p["value"] <= K.YEAR1:
                    two[owner[no]].add("on_time" if p["value"] <= due[no] else "late")
    check("Ask B D6: every council's on-time figure, and every late figure above nil, carries a receipt that paid a "
          "sundry invoice too (%s)" % {c: sorted(v) for c, v in two.items()},
          all(two[c] >= {k for k in ("on_time", "late") if B[c][k] > 0} for c in K.CODES))
    check("Ask B D6: no file says what a receipt line's amount is",
          not any(re.search(r"receipt'?s amount|amount of the receipt|amount on (each|every) line", flat(t), re.I)
                  for t in tx.values()))
    ye = [p for p in T["pays"] if p["value"] <= K.YEAR1 and p["post"] > O.CASHBOOK_TO
          and any(no in owner for no, _ in p["covers"])]
    ye_inst = sorted(owner[no] for p in ye for no, _ in p["covers"] if no in owner)
    check("Ask B D5: the cash book's year ends on 31 March by posting date, so the %d credits of 31 March paying eight "
          "councils' fourth instalments (%s) are in the bank's report only" % (len(ye), ", ".join(ye_inst)),
          len(ye) == 7 and len(ye_inst) == 8 and all(p["value"] == K.YEAR1 for p in ye)
          and not set(cb["Narrative"]) & {p["ref"] for p in ye} and all(key[(p["ref"], "%.2f" % p["amount"])] == 1
                                                                        for p in ye))
    levy_cb = cb[cb["Account"] == O.LEVY_ACCOUNT]
    prior = levy_cb[~levy_cb["Invoice"].isin(sales)]
    check("Ask B D5: the cash book's first levy postings are the 2024-25 fourth instalments under invoices the 2025-26 "
          "list does not hold, and the bank's report starts after them (%d)" % len(prior),
          len(prior) == len(O.PRIOR_Q4_ON_DUE) and not set(prior["Narrative"]) & set(bk["Reference"]))
    joint = bk[bk["Remitter"] == O.JOINT_PAYER]
    hnm4, wds4 = [[i for i in levy_inv if i["code"] == c and i["q"] == 4][0] for c in O.JOINT]
    check("Ask B D7: the shared finance service's one transfer, on 31 March, quotes Hornmere's fourth invoice only, at "
          "the two councils' fourth instalments together", len(joint) == 1 and
          joint["Reference"].item().startswith(hnm4["no"]) and
          abs(float(joint["Amount"].item()) - hnm4["amount"] - wds4["amount"]) < .005)
    fnc = bk[(bk["Remitter"] == O.BANK_NAME["FNC"]) & bk["Reference"].str.contains("INV26")]
    fnc_inv = [i for i in levy_inv if i["code"] == "FNC" and i["q"] == 4][0]
    grb_inv = [i for i in levy_inv if i["code"] == "GRB" and i["q"] == 4][0]
    check("Ask B D7: Ferncote's fourth payment quotes Grinbury's fourth invoice at Ferncote's own amount, in the bank's "
          "report only", len(fnc) == 1 and fnc["Reference"].item().startswith(grb_inv["no"]) and
          abs(float(fnc["Amount"].item()) - fnc_inv["amount"]) < .005 and fnc_inv["amount"] != grb_inv["amount"] and
          fnc["Reference"].item() not in set(cb["Narrative"]))
    own = bk[bk["Remitter"].isin([O.BANK_NAME[c] for c in O.REF_STYLE])]
    check("Ask B D7: Keldmere, Ravenwick and Chaldale quote their own references, never an invoice (%d credits)" % len(own),
          len(own) >= 11 and not own["Reference"].str.contains("INV2").any())
    hits = lambda val_: sum(not miss(val_[c][x], B[c][x]) for c in K.CODES for x in ("on_time", "late"))
    stopsB = {"natural (the cash book's posting dates, its lines added up)":
              (dict(dates="posted", lines="line", yearend="none"), 0),
              "the cracker (the bank's value dates, the cash book's lines added up, the year end's credits against the "
              "payers' open invoices)": (dict(dates="bank", lines="line", yearend="payer"), 2),
              "the bank's value dates, the cash book's lines added up, nothing past the cash book":
              (dict(dates="bank", lines="line", yearend="none"), 2),
              "each receipt once, nothing past the cash book": (dict(dates="bank", lines="receipt", yearend="none"), 10),
              "each receipt once, the year end's credits on the invoice each quotes":
              (dict(dates="bank", lines="receipt", yearend="reference"), 11),
              "the posting dates, everything else right": (dict(dates="posted", lines="receipt", yearend="payer"), 0)}
    for lab, (kw, k) in stopsB.items():
        h = hits(G.levy_timeliness(target, **kw))
        check("Ask B stop, %s: hits %d of eighteen" % (lab, h), h <= k)
    blind = sum(not miss(levy25[c], B[c]["on_time"]) + (not miss(0, B[c]["late"])) for c in K.CODES)
    check("Ask B blind response (every council paid every instalment on time) recovers %d of eighteen, under a quarter"
          % blind, blind <= 4)
    check("Ask A: the tariff notes state the VAT rate, in that file only", "standard rate of 20 per cent" in
          tx[K.FILES["tariffs"]] and sum("20 per cent" in flat(t) for t in tx.values()) == 1)
    check("SF11: the resolution states the apportionment on each council's number, rounded once at the end",
          "multiplied by its number and divided by the nine councils' numbers added together, rounded to the nearest "
          "pound" in papers)
    check("Ask B no figure is a multiple of 100 pounds or a share of the levy that reads as placed (every on-time figure "
          "differs from its council's levy but two)", sum(abs(B[c]["on_time"] - levy25[c]) < .005 for c in K.CODES) <= 2)

    # ------------------------------------------------------------------------------------------ Ask C
    truthC = collections.Counter()
    for (cp, h), n in T["ent"].items():
        truthC[K.CP_PARK[cp]] += n
    V = o["vehicles"]
    check("Ask C golden equals the vehicles that entered each park (%s)" % dict(V), all(V[p] == truthC[p]
                                                                                          for p in K.CHARGING))
    cn = pd.read_csv(G.f(target, "counts"))
    cn["hour"] = pd.to_datetime(cn["hour_start_utc"])
    model = dict(zip(reg["asset_id"], reg["model"]))
    check("Ask C D6: the Water Park's BX-2 loop reports from 18 August 2025 and its year's counts are even (in and out)",
          model["C-09"] == "BX-2" and cn.loc[cn["counter_id"] == "C-09", "hour"].min() == pd.Timestamp(O.SWAP_DW).floor("h")
          and int(cn.loc[cn["counter_id"] == "C-09", "vehicles"].sum()) % 2 == 0)
    tk_day = G.machines(target)
    dw_tix = tk_day[tk_day["park"] == "DW"].groupby("date").size()
    cday = cn.assign(day=G.local_date(cn["hour"].dt.strftime("%Y-%m-%dT%H:%M:%SZ")).values)
    per_tix = {cid: cday[cday["counter_id"] == cid].groupby("day")["vehicles"].sum().sum() /
              dw_tix[dw_tix.index.isin(cday.loc[cday["counter_id"] == cid, "day"])].sum() for cid in ("C-04", "C-09")}
    check("Ask C D6 (write-up step 7): C-09's counts per ticket at the Water Park's machines run at twice C-04's "
          "(%.2f against %.2f)" % (per_tix["C-09"], per_tix["C-04"]), 1.8 <= per_tix["C-09"] / per_tix["C-04"] <= 2.2)
    th_tix = tk_day[tk_day["machine_id"].isin(["PM-2401", "PM-2402"])].groupby("date").size()
    lv_tix = tk_day[tk_day["machine_id"] == "PM-2302"].groupby("date").size()
    c7 = cday[(cday["counter_id"] == "C-07") & (cday["hour"] < pd.Timestamp(O.SWAP_LM_TH).floor("h"))].groupby("day")[
        "vehicles"].sum()
    c7 = c7[c7 > 0]
    r_th = np.corrcoef(c7.values, th_tix.reindex(c7.index, fill_value=0).values)[0, 1]
    r_lv = np.corrcoef(c7.values, lv_tix.reindex(c7.index, fill_value=0).values)[0, 1]
    check("Ask C D2 (write-up step 7): before 14 October 2025 C-07's daily counts follow Thruscote Hall's machine "
          "payments (correlation %.2f) far more than Lockshaw's valley car park's (%.2f)" % (r_th, r_lv), r_th > r_lv + .2)
    a5 = cn[cn["counter_id"] == "C-05"].set_index("hour")["vehicles"]
    a10 = cn[cn["counter_id"] == "C-10"].set_index("hour")["vehicles"]
    both = a5.index.intersection(a10.index)
    check("Ask C D1: C-05 and C-10 both report Gorthfield's lakeside entrance for %d hours (3 to 24 February 2026), "
          "reading one loop" % len(both), 480 <= len(both) <= 520 and (a5[both] == a10[both]).all())
    c7 = cn[cn["counter_id"] == "C-07"]
    check("Ask C D2: C-07 reports all year; the register has it at Lockshaw's valley car park, beside C-08 (removed 14 "
          "October 2025), and Thruscote has no counter before C-11's install that day",
          c7["hour"].min() <= pd.Timestamp("2025-04-01") and reg.loc[reg["asset_id"] == "C-07", "site_code"].item() ==
          "LM-V" and cn.loc[cn["counter_id"] == "C-11", "hour"].min() == pd.Timestamp(O.SWAP_LM_TH).floor("h"))
    stopsC = {"natural (register sites, every row, passes as counted)": (G.vehicles(target, False, False, False), 4),
              "BX-2 halved, register sites, every row": (G.vehicles(target, True, False, False), 3),
              "halved, one counter per car park and hour, register sites": (G.vehicles(target, True, False, True), 1),
              "halved, the moved counter placed, every row": (G.vehicles(target, True, True, False), 1)}
    for lab, (val_, k) in stopsC.items():
        n = sum(miss(val_[p], V[p], 100) for p in K.CHARGING)
        check("Ask C stop, %s: misses %d of four by 100 vehicles or more" % (lab, n), n == k)

    # ------------------------------------------------------------------------------------------ cracker sheet
    import cracker as CR
    rows_cr = CR.score(target, o)
    _, lc, ls, pair = CR.ceiling(rows_cr, rows_cr)
    kept = {a: h for a, _n, h in rows_cr}
    nil = [c for c in K.CODES if B[c]["late"] == 0]
    check("cracker.py on the cut pack: the cracker keeps only the nil late figures every response files (%s; nil late: "
          "%s), pair ceiling %.1f" % (", ".join("%s %d of %d" % (a, h, n) for a, n, h in rows_cr), ", ".join(nil), pair),
          kept["A"] == 0 and kept["C"] == 0 and kept["B"] == len(nil) == 2 and pair <= 40)

    # ------------------------------------------------------------------------------------------ separation
    main = [K.FILES[k] for k in ("wave1", "wave2", "weights")]
    tmp = copy_without(target, [f for f in os.listdir(target) if f not in main])
    try:
        check("separation: the main call computes from the survey files alone, every ask file deleted",
              levies(G.residents(G.survey(tmp))) == L)
    finally:
        shutil.rmtree(tmp)
    check("separation: no device row in the main population (the survey and its weights carry no machine, counter, "
          "invoice or bank field)", not any(re.search(r"machine|counter|invoice|txn|ticket", c, re.I)
                                            for c in w1.columns.tolist() + w2.columns.tolist() + wt.columns.tolist()))
    base = all_figures(target)
    meta = json.load(open(os.path.join(os.path.dirname(target), "metadata.json")))
    for k in DL.DISTRACTORS:
        tmp = copy_without(target, [K.FILES[k]])
        try:
            check("distractor deleted, every figure unchanged: %s" % K.FILES[k], all_figures(tmp) == base)
        finally:
            shutil.rmtree(tmp)
    check("metadata names the distractors by bundle path", sorted(meta["distractor_files"]) ==
          sorted("target/" + K.FILES[k] for k in DL.DISTRACTORS))
    check("no shipped file name says distractor, decoy or answer",
          not any(re.search(r"distract|decoy|answer|truth", f, re.I) for f in os.listdir(target)))

    # ------------------------------------------------------------------------------------------ input gates
    files = os.listdir(target)
    check("input gate: 10 or more files (%d)" % len(files), len(files) >= 10)
    check("input gate: 3 or more formats (%s)" % sorted({f.rsplit(".", 1)[1] for f in files}),
          len({f.rsplit(".", 1)[1] for f in files}) >= 3)
    check("input gate: a file of 25,000 or more rows (machine payments, %d)" % len(mt), len(mt) >= 25000)
    check("input gate: two or more distractors", len(meta["distractor_files"]) >= 2)
    man = pd.read_csv(G.f(target, "manifest"))
    check("H9: the manifest lists every shipped file but itself", set(man["file"]) == set(files) - {K.FILES["manifest"]})
    pages = {f: "%d pages" % len(PdfReader(os.path.join(target, f)).pages) for f in files if f.endswith(".pdf")}
    check("H3: the manifest's page counts are the PDFs' (%s)" % pages,
          all(man.loc[man["file"] == f, "rows or pages"].item() == v for f, v in pages.items()))

    # ------------------------------------------------------------------------------------------ realism
    blob = " ".join(tx.values())
    check("realism: no em or en dash in any shipped text", "—" not in blob and "–" not in blob)
    check("realism: the only people named are the drawn personas", not re.search(
        r"\b(John|Jane) (Doe|Smith)\b|Acme|Company [A-Z]\b", blob) and K.PEOPLE["finance"] in papers and
        K.PEOPLE["contractor"] in note)
    check("realism: every shipped file carries the export's timestamp",
          len({int(os.path.getmtime(os.path.join(target, f))) for f in files}) == 1)
    piv = cn.pivot_table(index="hour", columns="counter_id", values="vehicles")
    same = []
    cols = list(piv.columns)
    for i, a in enumerate(cols):
        for b in cols[i + 1:]:
            j = piv[[a, b]].dropna()
            j = j[(j[a] > 0) | (j[b] > 0)]
            if len(j) and (j[a] == j[b]).mean() > .5:
                same.append((a, b))
    check("realism: no two counters repeat each other's hours but C-05 and C-10 on one loop", same == [("C-05", "C-10")])
    daily = mt.assign(day=mt["paid_at_utc"].str[:10]).groupby(["machine_id", "day"]).size()
    check("realism: no machine takes the same number of payments every day (most common daily count under 6 per cent "
          "of its days)", daily.groupby(level=0).apply(lambda s: s.value_counts().iloc[0] / len(s)).max() < .06)
    tk = mt.drop_duplicates(["machine_id", "ticket_no"]).sort_values(["machine_id", "paid_at_utc"])
    check("realism: each machine's ticket numbers rise with time", tk.groupby("machine_id")["ticket_no"].apply(
        lambda s: s.is_monotonic_increasing).all())
    check("realism: no round headline figure (the asks' figures and the levies are not multiples of 100)",
          not any(x % 100 == 0 for x in list(L.values()) + [int(np.floor(v + .5)) for v in inc.values()] +
                  [int(v) for v in V.values()]))
    perpark = d.groupby("park").size()
    check("realism: every park's interview total differs (%s)" % ", ".join("%s %d" % kv for kv in perpark.items()),
          perpark.nunique() == len(perpark))
    q9 = pd.concat([d["v_" + q] for q in K.PARK_CODES])
    check("realism: visit counts heap at round numbers as people give them (10, 20, 50, 100 more common than their "
          "neighbours)", all(int((q9 == k).sum()) > 2 * int((q9 == k + 1).sum()) for k in (10, 20, 50, 100)))
    check("realism: the residents' estimates are not round and differ from the populations' ordering",
          not any(x % 100 == 0 for x in [int(v) for v in N.values()]) and
          sorted(K.CODES, key=lambda c: -N[c]) != sorted(K.CODES, key=lambda c: -W.POP[c][2]))
    check("realism: every date in the pack is on or before the export (6 January 2027) and the levy papers'",
          max(pd.to_datetime(cb["Posting date"], format="%d/%m/%Y")).date() <= dt.date(2026, 4, 7) and
          pd.to_datetime(mt["paid_at_utc"]).max().date() <= dt.date(2026, 3, 31))
    for f in files:
        if f.endswith((".xlsx", ".docx")):
            with zipfile.ZipFile(os.path.join(target, f)) as z:
                core = z.read("docProps/core.xml").decode()
            check("H1 %s: no writer's name, a stamp in the fiction" % f, "openpyxl" not in core and "python" not in core
                  and re.search(r"<dcterms:modified[^>]*>202[5-7]-", core) is not None)
    for f in files:
        if f.endswith(".pdf"):
            raw = open(os.path.join(target, f), "rb").read()
            check("H1 %s: no ReportLab producer, dated in the fiction" % f, b"ReportLab" not in raw and
                  re.search(rb"/CreationDate \(D:202[5-7]", raw) is not None)
    check("H21 invented names clear of real UK places (no distinctive word within two edits of the near-miss register, "
          "no retired name surviving)", h21(blob))

    # ------------------------------------------------------------------------------------------ figures.json
    check("figures.json carries the call as the goldens state it", F["call"]["levy"][K.NAME["KDM"]] == L["KDM"] and
          F["call"]["valley_share"] == round(v["golden"], 1) and F["ask_c"][K.PARK_NAME["GF"]] == V["GF"])
    print("checks: %d assertions passed" % len(PASSED))
    return PASSED


# the real UK names nearest the invented words (places, NHS bodies and parks, each within three edits of an invented
# word), and the names this build retired (two after the screen, the rest for one spelling pattern across them all)
REAL = ['alnwick', 'appleshaw', 'arnside', 'askam', 'askern', 'askham', 'aylsham', 'balsham', 'banbury', 'banstead',
        'barnack', 'bassett', 'bearsden', 'bere', 'biggar', 'bishop', 'bishops', 'boosbeck', 'bosham', 'bramhope',
        'bramley', 'brandon', 'bratton', 'bredbury', 'brixworth', 'bulbeck', 'bumpstead', 'bunbury', 'carlisle',
        'carlton', 'carluke', 'carnock', 'caythorpe', 'chale', 'chalgrave', 'chalke', 'cheadle', 'cheswick',
        'childwall', 'chirnside', 'chiswick', 'clumber', 'codicote', 'consett', 'cookstown', 'corsham', 'cosham',
        'cowie', 'crowle', 'croxdale', 'croyde', 'crundale', 'culmore', 'danbury', 'deeside', 'denbury', 'denshaw',
        'dinmore', 'downside', 'dunvant', 'earnock', 'earsham', 'edgcott', 'edgworth', 'edworth', 'elham', 'ella',
        'elmet', 'elmham', 'elmstead', 'elstead', 'elstow', 'elsworth', 'eltham', 'ernest', 'evenwood', 'evesham',
        'exmoor', 'eynsham', 'fenwick', 'ferndale', 'ferndown', 'ferne', 'flamstead', 'flockton', 'foulsham',
        'fovant', 'fowlmere', 'fownhope', 'frogmore', 'garvagh', 'gisburn', 'gosfield', 'gransden', 'gringley',
        'grinstead', 'halsham', 'halstead', 'hampstead', 'hanslope', 'harlech', 'harpole', 'harrogate', 'hartfield',
        'heathfield', 'highgate', 'hillside', 'holbeck', 'hope', 'horncastle', 'hornsea', 'hornsey', 'horse',
        'horsham', 'hose', 'huncote', 'hunwick', 'husthwaite', 'instow', 'istead', 'ixworth', 'kellet', 'kelloe',
        'kelsey', 'kielder', 'kintbury', 'kinvere', 'larbert', 'larkfield', 'ledston', 'linthwaite', 'llanarth',
        'longhope', 'longside', 'mallet', 'margate', 'markyate', 'marshfield', 'martock', 'meldreth', 'melksham',
        'menston', 'mere', 'mersea', 'morcott', 'newbold', 'northaw', 'oakmere', 'ogmore', 'oldham', 'ongar',
        'orsett', 'ossett', 'ottershaw', 'oxenhope', 'padstow', 'palmers', 'peakirk', 'pelham', 'penarth', 'polruan',
        'pudsey', 'pulham', 'redwick', 'ridgmont', 'rigside', 'roche', 'rogiet', 'rossett', 'rowde', 'rushden',
        'rushmoor', 'ryhope', 'sapcote', 'selkirk', 'selsey', 'shaw', 'shepreth', 'sherwood', 'stanhope', 'sunbury',
        'talgarth', 'tangmere', 'tarbert', 'tenbury', 'thore', 'thorner', 'thorpe', 'tidbury', 'tilbury', 'tintern',
        'tisbury', 'topsham', 'tregarth', 'ushaw', 'walsham', 'wansbeck', 'wanstead', 'warwick', 'wedmore',
        'wendron', 'werneth', 'weston', 'wigmore', 'wilmcote', 'wilstead', 'windsor', 'wingate', 'winslow',
        'winston', 'wishaw', 'wistow', 'woodcote', 'woodside']
RETIRED = ["tarnwick", "velbrook", "kesvham", "tarvstead", "jeskley", "lorvborne", "glevby", "varvcote", "teskfield",
           "erskden", "cresvley", "corvgarth", "lesvley", "thorvcote", "fenvstead", "oskden"]
GENERIC = set("""north south east west upper lower great little mount port point river bay beach hill hills heights park
vale valley plain flat lake lakes springs island head bridge the of st saint new old central city council district borough
county town village regional region and moor moors forest wood woods common green country parks authority trust upon on
in le la by cum with under over end farm hall court lane road street way ltd co services research parking merchant
fencing tree grid energy plant hire print waste management builders ecology signs water security cafe farms masts
graziers upland uplands northern main lakeside summit""".split())


def h21(blob):
    names = {K.AUTH, K.CONTRACTOR, O.APP_PROVIDER, O.ACQUIRER} | {c[1] for c in K.COUNCILS} | \
        {p[1] for p in K.PARKS} | {v for vs in K.SETTLEMENTS.values() for v in vs} | set(O.SUPPLIERS) | set(O.PAYERS)
    words = {w for n in names for w in re.sub(r"[^a-z ]", " ", n.lower()).split() if w not in GENERIC and len(w) >= 4}
    for w in words:
        for r in REAL:
            if lev(w, r) <= 2 or (len(r) >= 5 and (w in r or r in w)):
                return False
    low = blob.lower()
    return not any(re.search(r"\b%s\b" % r, low) for r in RETIRED)
