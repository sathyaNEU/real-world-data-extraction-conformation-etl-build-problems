"""Independent verifier for task124: reads only target/ and golden/, on its own parsing path (csv module, the capital
plan's raw XML, the plan order from the plan's own table), recomputes the call, the default path and both asks, and
checks every owed figure in the golden workbook.

    python3 task124/generator/verify.py [target] [golden]"""
from __future__ import annotations

import csv
import datetime as dt
import math
import os
import re
import sys
import zipfile
from collections import Counter, defaultdict
from zoneinfo import ZoneInfo

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
EASTERN = ZoneInfo("America/New_York")


def rows(target, name):
    with open(os.path.join(target, name), newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def find(target, pattern):
    hits = [f for f in os.listdir(target) if re.fullmatch(pattern, f)]
    assert len(hits) == 1, (pattern, hits)
    return hits[0]


def era(y):
    return 0 if y < 1930 else 1 if y < 1950 else 2 if y < 1989 else 3


def plan_order(target):
    from pypdf import PdfReader
    text = "\n".join(p.extract_text() for p in PdfReader(os.path.join(target, find(target, r".*plan_rev2.*\.pdf")))
                     .pages)
    found = re.findall(r"^\s*(\d{1,2})\s+(RZ-\d\d)\b", text, re.M)
    order = [z for _, z in sorted(((int(n), z) for n, z in found))]
    assert len(order) == 12 and len(set(order)) == 12, found
    return order


def appropriation(target):
    with zipfile.ZipFile(os.path.join(target, find(target, r".*capital_improvement_plan.*\.docx"))) as z:
        xml = z.read("word/document.xml").decode()
    text = re.sub(r"<[^>]+>", "|", xml)
    m = re.search(r"Lead service line replacement, construction\|+([\d.]+)\|", text)
    return float(m.group(1)) * 1e6


def unit_costs(target):
    import openpyxl
    ws = openpyxl.load_workbook(os.path.join(target, find(target, r"unit_cost_schedule_\d{4}\.xlsx")),
                                data_only=True).worksheets[0]
    prices, mob = {}, None
    for r in ws.iter_rows(values_only=True):
        if isinstance(r[0], str) and re.fullmatch(r"RZ-\d\d", r[0]):
            prices[r[0]] = (float(r[2]), float(r[3]))
        if isinstance(r[0], str) and r[0].lower().startswith("mobilisation"):
            mob = float(r[2])
    return prices, mob


def main(target=None, golden=None):
    target = target or os.path.join(TASK, "target")
    golden = golden or os.path.join(TASK, "golden")
    inv = rows(target, find(target, r"hfrwd_service_line_inventory_\d{8}\.csv"))
    xw = {r["legacy_premise_no"]: r["sl_id"] for r in rows(target, "cis_gis_crosswalk_2023.csv")}
    code_cls = {r["result_code"]: r["inventory_class"] for r in rows(target, "sl_material_result_codes.csv")}
    visits = defaultdict(list)
    for r in rows(target, "sl_field_verifications_2021_2026.csv"):
        sl = r["sl_id"] or xw[r["premise_no"].rjust(8, "0")]
        d = (dt.datetime.strptime(r["verified_date"], "%m/%d/%Y").date() if "/" in r["verified_date"]
             else dt.date.fromisoformat(r["verified_date"]))
        visits[sl].append((d, code_cls[r["result_code"]]))
    seg_lines, seg_seen = Counter(), Counter()
    for r in inv:
        seg_lines[r["main_segment"]] += 1
        seg_seen[r["main_segment"]] += r["sl_id"] in visits
    whole = {s for s in seg_lines if seg_seen[s] == seg_lines[s]}
    # the class of each main, read from the tap cards on it: lead or galvanized if any card reads LEAD or GALV
    lg_main = defaultdict(bool)
    has_card = set()
    for r in inv:
        if r["tap_card_material"]:
            has_card.add(r["main_segment"])
            lg_main[r["main_segment"]] |= r["tap_card_material"] in ("LEAD", "GALV")
    klass = {m: ("LG" if lg_main[m] else "C") for m in has_card}
    n = {k: [0, 0, 0] for k in ("whole", "all")}
    h = {k: [0, 0, 0] for k in ("whole", "all")}
    cn, ch = Counter(), Counter()
    for r in inv:
        y = int(r["year_built"])
        if r["sl_id"] in visits and r["tap_card_material"] == "" and 0 < y < 1989:
            hit = sorted(visits[r["sl_id"]])[0][1] in ("LEAD", "GRR")
            for k in ("whole", "all"):
                if k == "all" or r["main_segment"] in whole:
                    n[k][era(y)] += 1
                    h[k][era(y)] += hit
            cn[(klass[r["main_segment"]], era(y))] += 1
            ch[(klass[r["main_segment"]], era(y))] += hit
    shares = {k: [h[k][e] / n[k][e] for e in range(3)] for k in n}
    cshare = {k: ch[k] / cn[k] for k in cn}
    pooled = {c: sum(ch[(c, e)] for e in range(3)) / sum(cn[(c, e)] for e in range(3)) for c in ("LG", "C")}
    # the share on lead or galvanized mains falls with the era, so the era-pooled rate misstates every era
    assert cshare[("LG", 0)] > cshare[("LG", 1)] > cshare[("LG", 2)] and \
        all(cshare[("C", e)] < 0.01 for e in range(3)), cshare
    order = plan_order(target)
    prices, mob = unit_costs(target)
    ap = appropriation(target)
    lgrr, unk = Counter(), defaultdict(lambda: [0, 0, 0])
    unk_cls = defaultdict(Counter)
    for r in inv:
        if r["sl_class"] in ("LEAD", "GRR"):
            lgrr[r["zone"]] += 1
        elif r["sl_class"] == "LEAD STATUS UNKNOWN":
            e = era(int(r["year_built"]))
            unk[r["zone"]][e] += 1
            unk_cls[e][klass[r["main_segment"]]] += 1
    rec_n, rec_h = [0, 0, 0], [0, 0, 0]
    for r in inv:
        y = int(r["year_built"])
        if r["tap_card_material"] and 0 < y < 1989:
            rec_n[era(y)] += 1
            rec_h[era(y)] += r["tap_card_material"] in ("LEAD", "GALV")
    rec_share = [rec_h[e] / rec_n[e] for e in range(3)]

    def schedule(p):
        out, cum = [], 0.0
        for z in order:
            exp = sum(unk[z][e] * p[e] for e in range(3))
            uf, ud = prices[z]
            cost = (lgrr[z] + exp) * uf + (sum(unk[z]) - exp) * ud + mob
            cum += cost
            out.append((z, lgrr[z] + exp, cost, cum, 2027 + math.ceil(cum / ap - 1e-9) - 1))
        return out
    # plan 4.3: the share expected among the district's Lead Status Unknown lines of each era. Each unknown line
    # carries the share found on the verified no-tap-card lines of mains of its class built in its era; the era's
    # share is their average
    left = [sum(unk_cls[e][k] * cshare[(k, e)] for k in ("LG", "C")) / sum(unk_cls[e].values()) for e in range(3)]
    flat = [sum(unk_cls[e][k] * pooled[k] for k in ("LG", "C")) / sum(unk_cls[e].values()) for e in range(3)]
    gold = schedule(left)
    whole_path = schedule(shares["whole"])
    dflt = schedule(shares["all"])
    card_path = schedule(rec_share)
    flat_path = schedule(flat)
    print("verify: %d segments verified whole; shares by class and era %s; unknown-line shares %s; every verification "
          "%s; whole segments %s; tap cards %s" % (
              len(whole), {"%s%d" % k: "%.4f" % v for k, v in sorted(cshare.items())}, ["%.4f" % x for x in left],
              ["%.4f" % x for x in shares["all"]], ["%.4f" % x for x in shares["whole"]],
              ["%.4f" % x for x in rec_share]))
    print("verify: last zone %s finished %d (tap cards %d, whole segments %d, default path %d); plan cost $%.1f million "
          "over %.2f programme years; one rate per kind across eras moves %d zones' lines" % (
              gold[-1][0], gold[-1][4], card_path[-1][4], whole_path[-1][4], dflt[-1][4], gold[-1][3] / 1e6,
              gold[-1][3] / ap, sum(1 for a, b in zip(gold, flat_path) if math.floor(a[1] / 10 + 0.5) !=
                                    math.floor(b[1] / 10 + 0.5))))

    # ask A
    comp = rows(target, "lsl_completions_2026.csv")
    latest = {}
    for r in sorted(comp, key=lambda r: (r["site_id"], r["closed_on"])):
        latest[r["site_id"]] = r
    import openpyxl
    ws = openpyxl.load_workbook(os.path.join(target, "change_orders_2026.xlsx"), data_only=True).worksheets[0]
    co = []
    head = None
    for r in ws.iter_rows(values_only=True):
        if r[0] == "CO no.":
            head = list(r)
        elif head and r[0]:
            d = dict(zip(head, r))
            co.append(([s.strip() for s in d["Sites covered"].split(",")], float(d["Amount ($)"])))
    tot, cnt = defaultdict(float), Counter()
    where = {}
    for sid, r in latest.items():
        t = dt.datetime.strptime(r["completed_utc"], "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=dt.timezone.utc)
        loc = t.astimezone(EASTERN).date()
        key = (r["contract_no"], (loc.month - 1) // 3 + 1)
        where[sid] = key
        tot[key] += float(r["paid_base"]) + float(r["paid_extra_ft"]) + float(r["paid_restoration"])
        cnt[key] += 1
    for sites, amt in co:
        keys = {where[s] for s in sites}
        assert len(keys) == 1
        tot[keys.pop()] += amt
    pay = {k: tot[k] / cnt[k] for k in cnt}

    # ask B: the agreement system's signed forms, and the lines the portal shows e-signed that the agreement system
    # holds no form for
    zone = {r["sl_id"]: r["zone"] for r in inv}
    agreed = defaultdict(set)
    held = set()
    for r in rows(target, "access_agreements_2022_2026.csv"):
        sl = r["sl_id"] or xw[r["premise_no"].rjust(8, "0")]
        held.add(sl)
        if r["status"] == "SIGNED":
            agreed[zone[sl]].add(sl)
    for r in rows(target, find(target, r"lslr_outreach_tracker_\d{8}\.csv")):
        if r["portal_status"] == "E-SIGNED" and r["sl_id"] not in held:
            agreed[zone[r["sl_id"]]].add(r["sl_id"])

    # the golden workbook
    gw = openpyxl.load_workbook(os.path.join(golden, "lslr_2027_programme.xlsx"), data_only=True)
    bad = []
    sched = {r[1]: r for r in gw["Zone schedule"].iter_rows(min_row=5, values_only=True) if r[1] in order}
    for z, repl, cost, cum, fy in gold:
        g = sched[z]
        if (g[3], g[4], g[6]) != (int(math.floor(repl / 10 + 0.5) * 10), int(math.floor(cost / 1e5 + 0.5) * 1e5), fy):
            bad.append(("zone", z, g[3:7], (repl, cost, fy)))
    for r in gw["2026 contract payments"].iter_rows(min_row=5, max_row=7, values_only=True):
        for q in (1, 2, 3, 4):
            if r[1 + q] != math.floor(pay[(r[0], q)] + 0.5):
                bad.append(("pay", r[0], q, r[1 + q], pay[(r[0], q)]))
    for r in gw["Access agreements"].iter_rows(min_row=5, values_only=True):
        if r[1] in order and r[3] != len(agreed[r[1]]):
            bad.append(("agree", r[1], r[3], len(agreed[r[1]])))
    assert not bad, bad
    assert gold[-1][4] == 2037 and min(card_path[-1][4], whole_path[-1][4], dflt[-1][4]) >= 2039, (
        gold[-1], card_path[-1], whole_path[-1], dflt[-1])
    assert sum(1 for a, b in zip(gold, flat_path) if math.floor(a[1] / 10 + 0.5) != math.floor(b[1] / 10 + 0.5)) >= 6
    print("verify: golden workbook agrees on 12 zones' lines, costs and finish years, 12 payments and 12 agreement "
          "counts")
    return gold, pay, agreed


if __name__ == "__main__":
    main(*(sys.argv[1:3]))
