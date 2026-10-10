"""Independent verifier for task124. Reads only the shipped bundle (target/), plus metadata.json for the distractor
gate, and recomputes every rung, rival-killer, calibration outcome and graded figure on its own code path.

    python3 task124/generator/verify_pack.py <target dir> [<metadata.json>] [--json out.json]
"""
from __future__ import annotations

import json
import math
import os
import re
import sys
from datetime import date, timedelta

import numpy as np
import pandas as pd

BOOKS = ["COAST", "EAST", "FWEST", "NORTH", "NCENT", "SCENT", "SOUTH", "WEST"]
NAMES = ["Coast", "East", "Far West", "North", "North Central", "South Central", "Southern", "West"]
EXPECT = {"answer": (120, 30, 0, 0, 210, 25, 15, 0), "R0": (115, 65, 35, 15, 40, 45, 60, 25),
          "R1": (110, 20, 0, 0, 255, 10, 5, 0), "R2": (135, 45, 0, 0, 145, 40, 35, 0),
          "R3": (130, 40, 0, 0, 165, 35, 30, 0)}
REFRIG = {"493120", "312113"}
RESULTS = []


def check(name, ok, detail=""):
    RESULTS.append((name, bool(ok), str(detail)[:300]))
    print(("PASS " if ok else "FAIL ") + name + ("" if ok else f"  [{detail}]"))


def pct90(x, kind="inclusive"):
    x = sorted(x)
    n = len(x)
    if kind == "inclusive":
        r = 0.9 * (n - 1)
    elif kind == "exclusive":
        r = 0.9 * (n + 1) - 1
    else:  # nearest rank
        return x[math.ceil(0.9 * n) - 1]
    lo = int(math.floor(r))
    return x[lo] + (r - lo) * (x[min(lo + 1, n - 1)] - x[lo])


def level(E, order=BOOKS, total=400, lot=5):
    a = {b: 0 for b in E}
    for _ in range(total // lot):
        best = None
        for b in order:
            if best is None or E[b] - a[b] > E[best] - a[best] + 1e-12:
                best = b
        a[best] += lot
    return a


def vec(a):
    return tuple(a[b] for b in BOOKS)


class Pack:
    def __init__(self, d):
        self.d = d
        rd = lambda f, **k: pd.read_csv(os.path.join(d, f), dtype=str, keep_default_na=False, **k)  # noqa: E731
        self.en = rd("enrollment_extract_20270409.csv")
        self.en["md"] = self.en["max_demand_kw"].astype(float)
        self.en["s"] = pd.to_datetime(self.en["start_date"]).dt.date
        self.en["e"] = [date.fromisoformat(x) if x else None for x in self.en["end_date"]]
        self.en["so"] = [x.toordinal() for x in self.en["s"]]
        self.en["eo"] = [x.toordinal() if x else 10 ** 9 for x in self.en["e"]]
        self.prem = rd("premise_register_20270409.csv").set_index("esi_id")
        self.pk = rd("ercot_summer_system_peaks.csv")
        self.peaks = {int(r.summer): (date.fromisoformat(r.peak_date), int(r.hour_ending)) for r in self.pk.itertuples()}
        st = pd.read_csv(os.path.join(d, "zone_settled_load_s17_s26.csv"))
        st["day"] = pd.to_datetime(st["oper_day"]).dt.date
        self.st = st.set_index(["day", "hour_ending"])
        self.reads = pd.read_parquet(os.path.join(d, "idr_hourly_reads_summers_2017_2026.parquet"))
        self.reads["read_date"] = pd.to_datetime(self.reads["read_date"]).dt.date
        self.cr = rd("bsaver_credits_2017_2026.csv")
        self.cr["day"] = pd.to_datetime(self.cr["event_date"]).dt.date
        self.acc = rd("billing_accounts.csv")
        self.years = sorted(self.peaks)

    def active(self, d):
        """Rows serving on day d, an AMEND row in service superseding the row it amends."""
        e, o = self.en, d.toordinal()
        a = e[(e["so"] <= o) & (e["eo"] >= o)]
        am = set(a.loc[a["record_type"] == "AMEND", "esi_id"])
        return a[(a["record_type"] == "AMEND") | ~a["esi_id"].isin(am)]

    def md(self, d, rows=False):
        e, o = self.en, d.toordinal()
        a = e[(e["so"] <= o) & (e["eo"] >= o)] if rows else self.active(d)
        return a.groupby("book")["md"].sum().reindex(BOOKS).fillna(0).to_dict()

    def load(self, d, he):
        return {b: float(self.st.loc[(d, he), b]) for b in BOOKS}


def main_call(P):
    yrs = P.years
    ext = date(2027, 4, 9)
    pk26 = P.peaks[2026][0]
    mdp = {y: P.md(P.peaks[y][0]) for y in yrs}
    lp = {y: P.load(*P.peaks[y]) for y in yrs}
    pooled = {b: {y: lp[y][b] / (mdp[y][b] / 1000.0) for y in yrs} for b in BOOKS}
    # members: accounts credited by the programme, to ESI IDs through billing accounts
    accts = set(P.cr["account_no"])
    mem = sorted(set(P.acc.loc[P.acc["account_no"].isin(accts), "esi_id"]))
    check("V01 fourteen member sites reached through credits and billing accounts, all refrigerated NAICS",
          len(mem) == 14 and all(P.prem.at[e, "naics_code"] in REFRIG for e in mem), len(mem))
    book27 = P.active(ext)
    rows27 = P.md(ext, rows=True)
    ded27 = book27.groupby("book")["md"].sum().reindex(BOOKS).to_dict()
    refrig27 = book27[book27["esi_id"].map(lambda e: P.prem.at[e, "naics_code"] in REFRIG)]
    read_esi = set(P.reads["esi_id"].unique())
    cen = refrig27[~refrig27["esi_id"].isin(read_esi)]
    check("V02 the new refrigerated premises: 31 in North Central, 186.0 MW, none with a summer read",
          len(cen) == 31 and set(cen["book"]) == {"NCENT"} and abs(cen["md"].sum() - 186000.0) < 0.05,
          (len(cen), cen["md"].sum()))
    centres = cen["md"].sum() / 1000.0
    am = P.en[P.en["record_type"] == "AMEND"]
    check("V03 the AMEND rows are all North Central, entered March 2027, same maximum demand as the row they amend",
          set(am["book"]) == {"NCENT"} and all(x.startswith("2027-03") for x in am["entered_on"])
          and (am.set_index("esi_id")["md"] == P.en[(P.en["record_type"] == "NEW") & P.en["esi_id"].isin(am["esi_id"])
                                              & (P.en["end_date"] == "")].set_index("esi_id")["md"]
               .reindex(am["esi_id"])).all(), len(am))
    # members' draw at the closed peaks (class factor) and on uncalled weekdays
    r = P.reads[P.reads["esi_id"].isin(mem)].copy()
    mdm = book27.set_index("esi_id")["md"].reindex(mem)
    r["md"] = r["esi_id"].map(mdm)
    r["y"] = [d.year for d in r["read_date"]]
    called = {}
    for a, d in zip(P.cr["account_no"], P.cr["day"]):
        called.setdefault(d.year, set()).add(d)
    r["called"] = [d in called.get(d.year, set()) for d in r["read_date"]]
    r["pkh"] = r["y"].map(lambda y: P.peaks[y][1])
    atpk = r[[(d == P.peaks[y][0] and h == P.peaks[y][1]) for d, y, h in zip(r["read_date"], r["y"], r["hour_ending"])]]
    cf = atpk["kwh"].sum() / atpk["md"].sum()
    unc = r[(r["hour_ending"] == r["pkh"]) & ~r["called"]]
    u = unc["kwh"].sum() / unc["md"].sum()
    check("V04 every closed system peak day is a credited (called) day", all(P.peaks[y][0] in called[y] for y in yrs))
    check("V05 class factor at the closed peaks about 0.60, uncalled draw about 0.90",
          0.595 <= cf <= 0.605 and 0.895 <= u <= 0.905, (round(cf, 5), round(u, 5)))
    share = max(sum(mdm[e] for e in mem if P.active(P.peaks[y][0])["esi_id"].eq(e).any() and
                    book27.set_index("esi_id").at[e, "book"] == b) / mdp[y][b] for b in BOOKS for y in yrs)
    check("V06 members under 1 per cent of every book's maximum demand in every closed summer", share < 0.01, share)

    def expo(base, extra=0.0, kind="inclusive"):
        E = {}
        for b in BOOKS:
            L = pct90([pooled[b][y] * base[b] / 1000.0 for y in yrs], kind)
            E[b] = L + (extra if b == "NCENT" else 0.0)
        return E
    hedges = hedge_mw(P.d)
    nc_no_c = dict(ded27, NCENT=ded27["NCENT"] - centres * 1000.0)
    R = {"R0": {b: pct90([lp[y][b] for y in yrs]) for b in BOOKS}, "R1": expo(rows27), "R2": expo(ded27),
         "R3": expo(nc_no_c, cf * centres), "R4": expo(nc_no_c, u * centres)}
    E = {k: {b: v[b] - hedges[b] for b in BOOKS} for k, v in R.items()}
    L = {k: level(v) for k, v in E.items()}
    check("V07 the answer recomputes: 120 / 30 / 0 / 0 / 210 / 25 / 15 / 0", vec(L["R4"]) == EXPECT["answer"], vec(L["R4"]))
    for k in ("R0", "R1", "R2", "R3"):
        check(f"V08 rung {k[1]} recomputes to {EXPECT[k]}", vec(L[k]) == EXPECT[k], vec(L[k]))
    post = {b: E["R4"][b] - L["R4"][b] for b in BOOKS}
    taken = [E["R4"][b] - 5 * k for b in BOOKS for k in range(L["R4"][b] // 5)]
    check("V09 lot decisions clear of a tie by 0.9 MW and post-block exposures clear of half-MW edges by 0.25 MW",
          min(taken) - max(post.values()) >= 0.9 and min(abs(v % 1 - 0.5) for v in post.values()) >= 0.25,
          (round(min(taken), 3), round(max(post.values()), 3)))
    check("V10 the next lot would have gone to Southern", max(post, key=post.get) == "SOUTH", post)
    # corridor on the centres' factor
    ok_f = [f / 10000 for f in range(8700, 9301) if vec(level({b: expo(nc_no_c, f / 10000 * centres)[b] - hedges[b]
                                                                  for b in BOOKS})) == EXPECT["answer"]]
    check("V11 corridor: every centre factor from 0.885 to 0.916 files the answer; 0.88 and 0.92 do not",
          min(ok_f) <= 0.885 and max(ok_f) >= 0.916 and 0.88 not in ok_f and 0.92 not in ok_f, (min(ok_f), max(ok_f)))
    for kind in ("exclusive", "nearest"):
        e2 = {b: expo(nc_no_c, u * centres, kind)[b] - hedges[b] for b in BOOKS}
        check(f"V12 P90 {kind} converges on the answer", vec(level(e2)) == EXPECT["answer"])
    ests = estimators(P, r, mem, called, mdm)
    check("V13 twelve estimators of the uncalled draw all inside [0.894, 0.906] and inside the corridor",
          len(ests) == 12 and all(0.894 <= v <= 0.906 and min(ok_f) < v < max(ok_f) for v in ests.values()),
          {k: round(v, 4) for k, v in ests.items()})
    return {"pooled": pooled, "mdp": mdp, "lp": lp, "mem": mem, "called": called, "u": u, "cf": cf, "E": E, "L": L,
            "post": post, "centres": centres, "book27": book27, "hedges": hedges, "r": r, "ests": ests,
            "corridor": (min(ok_f), max(ok_f))}


def estimators(P, r, mem, called, mdm):
    m = r[(r["hour_ending"] == r["pkh"]) & ~r["called"]].copy()
    m["s"] = m["kwh"] / m["md"]
    per_y = {y: g["kwh"].sum() / g["md"].sum() for y, g in m.groupby("y")}
    per_site = m.groupby("esi_id")["s"].mean()
    hot = hot_days(P, called)[0]
    mh = m[[d in hot for d in m["read_date"]]]
    # programme baseline at the peak hour on each peak day: ten most recent uncalled business days
    num = den = 0.0
    for y in P.years:
        pkd, he = P.peaks[y]
        hol = summer_holidays(y)
        prior = sorted(d for d in set(r.loc[r["y"] == y, "read_date"]) if d < pkd and d.weekday() < 5 and d not in hol
                       and d not in called[y])
        for e in mem:
            s = r[(r["esi_id"] == e) & (r["hour_ending"] == he) & r["read_date"].isin(prior)].sort_values("read_date")
            if len(s) == 0 or not ((r["esi_id"] == e) & (r["read_date"] == pkd)).any():
                continue
            num += s["kwh"].tail(10).mean()
            den += mdm[e]
    return {"pooled MD-weighted": m["kwh"].sum() / m["md"].sum(), "pooled mean ratio": m["s"].mean(),
            "pooled median ratio": m["s"].median(), "by summer mean": float(np.mean(list(per_y.values()))),
            "by summer median": float(np.median(list(per_y.values()))), "by summer max": max(per_y.values()),
            "by summer min": min(per_y.values()),
            "by site MD-weighted": float((per_site * mdm[per_site.index]).sum() / mdm[per_site.index].sum()),
            "by site mean": per_site.mean(), "2026 only": per_y[2026],
            "hottest-decile uncalled": mh["kwh"].sum() / mh["md"].sum(), "programme baseline": num / den}


def summer_holidays(y):
    j = date(y, 7, 4)
    j = j - timedelta(days=1) if j.weekday() == 5 else j + timedelta(days=1) if j.weekday() == 6 else j
    l = date(y, 9, 1)
    while l.weekday() != 0:
        l += timedelta(days=1)
    return {j, l}


def hot_days(P, called):
    tot = P.st[BOOKS].sum(axis=1).groupby(level=0).max()
    out, counts = set(), {}
    for y in P.years:
        t = tot[[d.year == y and d.weekday() < 5 for d in tot.index]]
        top = t.sort_values(ascending=False).index[:math.ceil(0.1 * len(t))]
        unc = [d for d in top if d not in called[y]]
        counts[y] = len(unc)
        out |= set(unc)
    return out, counts


def hedge_mw(d):
    from openpyxl import load_workbook
    ws = load_workbook(os.path.join(d, "hedge_positions_20270409.xlsx"), read_only=True)["Summer 2027"]
    out = {}
    for row in ws.iter_rows(values_only=True):
        if row and isinstance(row[0], str) and row[0].endswith(" total"):
            vals = [v for v in row[3:7]]
            assert len(set(vals)) == 1, row
            out[row[0].split()[0]] = int(vals[0])
    return out


def report_tables(d):
    from pypdf import PdfReader
    txt = "\n".join(p.extract_text() for p in PdfReader(os.path.join(d, "summer_2026_risk_report.pdf")).pages)
    lines = [x.strip() for x in txt.splitlines() if x.strip()]
    t1 = lines[lines.index("Table 1. Summer 2026 at the system peak hour.") - 24: lines.index("Table 1. Summer 2026 at the system peak hour.")]
    i = next(k for k, x in enumerate(lines) if x.startswith("Table 2."))
    t2 = lines[i - 88: i]
    tab1 = {NAMES[k]: (float(t1[3 * k + 1].replace(",", "")), float(t1[3 * k + 2].replace(",", ""))) for k in range(8)}
    tab2 = {}
    for k in range(8):
        seg = t2[11 * k: 11 * k + 11]
        assert seg[0] == NAMES[k], seg
        tab2[NAMES[k]] = [int(x.replace(",", "")) for x in seg[1:]]
    return tab1, tab2


def corpus(P, M):
    yrs = P.years
    pk26 = P.peaks[2026][0]
    tab1, tab2 = report_tables(P.d)
    mine = {b: [M["pooled"][b][y] * M["mdp"][2026][b] / 1000.0 for y in yrs] for b in BOOKS}
    hits = sum(int(math.floor(mine[b][j] + 0.5)) == tab2[n][j] for b, n in zip(BOOKS, NAMES) for j in range(10))
    check("V14 the per-kW replay reproduces all 80 cells of the close-out's replay table", hits == 80, hits)
    ok1 = all(abs(tab1[n][0] - M["mdp"][2026][b] / 1000.0) < 0.051 and abs(tab1[n][1] - M["lp"][2026][b]) < 0.051
              for b, n in zip(BOOKS, NAMES))
    check("V15 Table 1 of the close-out recomputes from the enrolment extract and the settled loads", ok1)
    edge = min(abs(v % 1 - 0.5) for b in BOOKS for v in mine[b])
    check("V16 every replay cell is at least 0.05 MW from a half-MW edge", edge >= 0.05, edge)

    def score(fn):
        return sum(int(math.floor(fn(b, y) + 0.5)) == tab2[n][j] for b, n in zip(BOOKS, NAMES) for j, y in enumerate(yrs))
    mdy = {y: P.md(date(y, 12, 31)) for y in yrs}
    mdj = {y: P.md(date(y, 6, 1)) for y in yrs}
    riv = {"settled unscaled": score(lambda b, y: M["lp"][y][b]),
           "divisor at year end": score(lambda b, y: M["lp"][y][b] / mdy[y][b] * M["mdp"][2026][b]),
           "divisor at 1 June": score(lambda b, y: M["lp"][y][b] / mdj[y][b] * M["mdp"][2026][b]),
           "rows not superseded": score(lambda b, y: M["lp"][y][b] / P.md(P.peaks[y][0], rows=True)[b] * P.md(pk26, rows=True)[b])}
    zmax = P.st[BOOKS].copy()
    own = {}
    for y in yrs:
        z = zmax[[d.year == y for d in zmax.index.get_level_values(0)]]
        for b in BOOKS:
            k = z[b].idxmax()
            own[(b, y)] = (float(z[b].max()), k[0])
    riv["zone's own peak hour"] = score(lambda b, y: own[(b, y)][0] / P.md(own[(b, y)][1])[b] * M["mdp"][2026][b])
    check("V17 every rival basis misses replay cells; settled loads unscaled match only 2026's eight",
          riv["settled unscaled"] == 8 and all(v < 80 for k, v in riv.items() if k != "rows not superseded"), riv)
    check("V18 the replay table is blind to the March 2027 amendments and the centres (no AMEND row or centre in a closed book)",
          riv["rows not superseded"] == 80)
    return {"rivals": riv}


def twins_and_credits(P, M):
    e = P.en[(P.en["record_type"] == "NEW") & (P.en["book"] == "NCENT") & (P.en["md"] == 2400.0)]
    cols = ["tdsp", "plan_code", "deposit_class", "meter_type", "premise_age_band", "broker_code", "start_date", "end_date"]
    pairs = [(a, b) for a in e.itertuples() for b in e.itertuples() if a.esi_id < b.esi_id
             and all(getattr(a, c) == getattr(b, c) for c in cols)]
    ok = len(pairs) == 1
    d, h = P.peaks[2026]
    vals = {}
    if ok:
        for x in pairs[0]:
            v = P.reads[(P.reads["esi_id"] == x.esi_id) & (P.reads["read_date"] == d) & (P.reads["hour_ending"] == h)]
            vals[x.esi_id] = (float(v["kwh"].iloc[0]), x.esi_id in M["mem"], P.prem.at[x.esi_id, "naics_code"])
    hi = max(v[0] for v in vals.values()) if vals else 0
    lo = min(v[0] for v in vals.values()) if vals else 1
    check("V19 twin pair: identical enrolment columns, the member drew 1.5x the warehouse at the 2026 peak",
          ok and hi / lo >= 1.5 and any(v[1] for v in vals.values()), vals)
    # credits recompute from the reads under the terms' baseline rule
    acc = P.acc[P.acc["esi_id"].isin(M["mem"])].groupby("account_no")["esi_id"].apply(list).to_dict()
    starts = P.en[P.en["esi_id"].isin(M["mem"])].set_index("esi_id")["s"]
    rr = P.reads[P.reads["esi_id"].isin(M["mem"])].set_index(["esi_id", "read_date", "hour_ending"])["kwh"]
    bad = 0
    pool = {}
    for esi in M["mem"]:
        ds = sorted(set(rr.loc[esi].index.get_level_values(0)))
        for y in P.years:
            hol = summer_holidays(y)
            pool[(esi, y)] = [x for x in ds if x.year == y and x.weekday() < 5 and x not in hol and x not in M["called"][y]]
    for row in P.cr.itertuples():
        dd = row.day
        y = dd.year
        tot = 0.0
        for esi in acc[row.account_no]:
            if starts[esi] > dd:
                continue
            days = [x for x in pool[(esi, y)] if x < dd][-10:]
            for hh in (15, 16, 17, 18):
                base = np.mean([rr[(esi, x, hh)] for x in days])
                tot += max(0.0, base - rr[(esi, dd, hh)])
        if int(round(tot)) != int(row.credited_kwh):
            bad += 1
            print("   credit mismatch", row.account_no, dd, round(tot, 4), row.credited_kwh)
    check("V20 every credit line recomputes from the reads under the baseline rule (ten most recent uncalled business days)",
          bad == 0, bad)
    hot, counts = hot_days(P, M["called"])
    check("V21 at least four weekdays of every summer's hottest decile carry no call", min(counts.values()) >= 4, counts)
    accts_nonref = set(P.cr["account_no"]) - set(acc)
    check("V22 no credit to an account without a member site", not accts_nonref, accts_nonref)


def asks(d):
    from docx import Document
    from openpyxl import load_workbook
    rd = lambda f: pd.read_csv(os.path.join(d, f), dtype=str, keep_default_na=False)  # noqa: E731
    proc = "\n".join(p.text for p in Document(os.path.join(d, "desk_procedures.docx")).paragraphs)
    tab = Document(os.path.join(d, "desk_procedures.docx")).tables[0]
    conv = {r.cells[0].text: re.search(r"(5x16|7x16)", r.cells[1].text).group(1) for r in tab.rows[1:]}
    check("V23 the procedures carry three approved brokers with their notional conventions and the conversion, "
          "decision, approval, version-of-record and MWh rules", len(conv) == 3 and all(x in proc for x in (
              "hours of the notional", "lowest accepted quote", "approved counterparty on the day",
              "latest amendment matched", "MWh in the period")), conv)
    cal = rd("trading_calendar_2027.csv")
    hol = {date.fromisoformat(x) for x in cal.loc[cal["nerc_holiday"] == "Y", "date"]}

    def hrs(shape, m, holidays=True):
        ds = [date(2027, m, 1) + timedelta(days=k) for k in range(31) if (date(2027, m, 1) + timedelta(days=k)).month == m]
        if shape == "7x16":
            return 16 * len(ds)
        return 16 * sum(1 for x in ds if x.weekday() < 5 and (not holidays or x not in hol))
    cp = rd("desk_counterparties.csv")
    appr = [(r.cpty_code, r.legal_name, date.fromisoformat(r.approved_from),
             date.fromisoformat(r.approved_to) if r.approved_to else None) for r in cp.itertuples()]
    q = rd("option_quotes_s27.csv")
    brk = {"Gulfline": "Gulfline Energy Brokers", "Trinity Basin": "Trinity Basin Capital Markets"}
    rows = [dict(ref=r.quote_id, rev=int(r.revision), broker=brk[r.broker], code=r.cpty_code, name=None,
                 zone=r.load_zone, m=int(r.delivery_month[-2:]), price=float(r.premium),
                 sent=date.fromisoformat(r.sent_at[:10])) for r in q.itertuples()]
    ws = load_workbook(os.path.join(d, "pecos_quote_sheet_apr2027.xlsx"), read_only=True)["Offers"]
    zmap = {"Houston LZ": "LZ_HOUSTON", "North LZ": "LZ_NORTH", "South LZ": "LZ_SOUTH", "West LZ": "LZ_WEST"}
    mon = {"Jun-27": 6, "Jul-27": 7, "Aug-27": 8, "Sep-27": 9}
    for row in list(ws.iter_rows(values_only=True))[4:]:
        if row[0]:
            sent = pd.to_datetime(row[6], format="%m/%d/%Y %I:%M %p").date()
            rows.append(dict(ref=row[0], rev=1, broker="Pecos Power Brokerage", code=None, name=row[1], zone=zmap[row[2]],
                             m=mon[row[3]], price=float(row[5]), sent=sent))
    Q = pd.DataFrame(rows)
    dec = rd("quote_decisions.csv").set_index(["quote_ref", "revision"])["decision"]
    Q["decision"] = [dec[(str(a), str(b))] for a, b in zip(Q["ref"], Q["rev"])]

    def approved(code, name, when):
        return any((c == code or (code is None and n == name)) and f <= when and (t is None or when <= t)
                   for c, n, f, t in appr)
    Q["ok"] = [approved(c, n, s) for c, n, s in zip(Q["code"], Q["name"], Q["sent"])]
    Q["reissued"] = [(c in ("C0231",)) or (n == "Saltgrass Energy LP") for c, n in zip(Q["code"], Q["name"])]
    bm = rd("book_zone_map.csv")

    def zone_of(book, m, latest_map=True):
        first = date(2027, m, 1)
        r = bm[bm["book"] == book]
        if not latest_map:
            r = r[r["effective_from"] <= "2026-12-31"]
            return r.sort_values("effective_from")["load_zone"].iloc[-1]
        r = r[(r["effective_from"] <= first.isoformat()) & ((r["effective_to"] == "") | (r["effective_to"] >= first.isoformat()))]
        assert len(r) == 1
        return r["load_zone"].iloc[0]

    def prem(holidays=True, decisions=True, approval=True, pecos=True, notional=True, latest=False, drop=False, map27=True):
        x = Q.copy()
        if not pecos:
            x = x[x["broker"] != "Pecos Power Brokerage"]
        if latest:
            x = x.sort_values("rev").groupby("ref").tail(1)
        if decisions:
            x = x[x["decision"] == "ACCEPTED"]
        if approval:
            x = x[x["ok"]]
        if drop:
            revd = set(Q.loc[Q["rev"] > 1, "ref"])
            x = x[~x["ref"].isin(revd) & ~x["reissued"]]
        x = x.assign(usd=[p * hrs(conv[b] if notional else "5x16", m, holidays) for p, b, m in zip(x["price"], x["broker"], x["m"])])
        best = x.groupby(["zone", "m"])["usd"].min()
        return {(b, m): float(best[(zone_of(b, m, map27), m)]) for b in BOOKS for m in (6, 7, 8, 9)}
    ans = prem()
    st = {"S1": prem(holidays=False, decisions=False, approval=False, pecos=False, notional=False, latest=True, map27=False),
          "S2": prem(holidays=False), "S3": prem(drop=True), "S4": prem(map27=False)}
    mv = {k: [c for c in ans if abs(v[c] - ans[c]) > 0.004] for k, v in st.items()}
    check("V24 premium stops: the natural path wrong on 24 or more of 32 cells, every stop at least 1 per cent off where "
          "it moves, the 2026 map moving only East", len(mv["S1"]) >= 24 and all(
              abs(st[k][c] / ans[c] - 1) >= 0.01 for k in st for c in mv[k]) and {c[0] for c in mv["S4"]} == {"EAST"},
          {k: len(v) for k, v in mv.items()})
    check("V25 every premium sits at least 5 cents from a half-dollar edge", min(abs(v % 1 - 0.5) for v in ans.values()) >= 0.05)
    # hedges
    bl = rd("trade_blotter_s27.csv")
    ml = rd("confirm_match_log.csv")
    pf = rd("portfolio_books.csv").set_index("portfolio")["book"]
    last = ml.sort_values("status_date", kind="mergesort").groupby(["trade_id", "amendment_no"]).tail(1)
    matched = set(zip(last.loc[last["status"] == "MATCHED", "trade_id"], last.loc[last["status"] == "MATCHED", "amendment_no"]))
    s = bl[(bl["delivery_start"] == "2027-06-01") & (bl["delivery_end"] == "2027-09-30")].copy()
    s["amend"] = s["amendment_no"].astype(int)
    s["book"] = s["portfolio"].map(pf)
    s["shape"] = s["product_code"].str[-4:].str.lower()
    s["price"] = s["fixed_price"].astype(float)
    s["mw"] = s["mw"].astype(float)

    def hp(version="matched", weight="mwh", holidays=True):
        x = s
        if version == "matched":
            x = x[[(t, a) in matched for t, a in zip(x["trade_id"], x["amendment_no"])]]
        elif version == "original":
            x = x[x["amend"] == 0]
        x = x.sort_values("amend").groupby("trade_id").tail(1)
        w = x["mw"] if weight == "mw" else x["mw"] * x["shape"].map(lambda sh: sum(hrs(sh, m, holidays) for m in (6, 7, 8, 9)))
        x = x.assign(w=w)
        return {b: float((g["price"] * g["w"]).sum() / g["w"].sum()) for b, g in x.groupby("book")}
    H = hp()
    hs = {"S1": hp("latest", "mw"), "S2": hp(weight="mw"), "S3": hp(holidays=False), "S4": hp("original")}
    check("V26 hedge price stops each at least $0.05/MWh off on six or more books, answers clear of half-cent edges",
          all(sum(abs(v[b] - H[b]) >= 0.05 for b in BOOKS) >= 6 for v in hs.values())
          and min(abs((x * 100) % 1 - 0.5) for x in H.values()) >= 0.15, {k: sum(abs(v[b] - H[b]) >= 0.05 for b in BOOKS) for k, v in hs.items()})
    mw = s[s["amend"] == 0].groupby("book")["mw"].sum().to_dict()
    check("V27 blotter MW per book ties the position report and amendments are price-only",
          mw == {k: float(v) for k, v in hedge_mw(d).items()} and s.groupby("trade_id")["mw"].nunique().max() == 1)
    return {"premium": {f"{b}|{m}": round(v, 2) for (b, m), v in ans.items()}, "hedge_price": {b: round(v, 2) for b, v in H.items()}}


def grid_and_lenders(P, M):
    from openpyxl import load_workbook
    yrs, pooled, hedges = P.years, M["pooled"], M["hedges"]
    book27, centres, u, cf = M["book27"], M["centres"], M["u"], M["cf"]
    ded27 = book27.groupby("book")["md"].sum().reindex(BOOKS).to_dict()
    nc0 = dict(ded27, NCENT=ded27["NCENT"] - centres * 1000.0)
    in26 = set(P.active(P.peaks[2026][0])["esi_id"])
    newgen = book27[(book27["book"] == "NCENT") & ~book27["esi_id"].isin(in26)
                    & ~book27["esi_id"].map(lambda e: P.prem.at[e, "naics_code"] in REFRIG)]["md"].sum()
    mem_md = book27[book27["esi_id"].isin(M["mem"])].groupby("book")["md"].sum().reindex(BOOKS).fillna(0).to_dict()

    def E(base, extra, plus=None):
        out = {}
        for b in BOOKS:
            L = pct90([pooled[b][y] * base[b] / 1000.0 for y in yrs]) + (extra if b == "NCENT" else 0.0)
            out[b] = L + (plus or {}).get(b, 0.0) - hedges[b]
        return out
    nc1 = dict(nc0, NCENT=nc0["NCENT"] - newgen)
    cells = {
        "partial: class factor on every new NC premise": E(nc1, cf * (centres + newgen / 1000.0)),
        "partial: uncalled draw on every new NC premise": E(nc1, u * (centres + newgen / 1000.0)),
        "members uncalled in 2027 too": E(nc0, u * centres, {b: (u - cf) * mem_md[b] / 1000.0 for b in BOOKS}),
    }
    lv = {k: vec(level(v)) for k, v in cells.items()}
    nc = {k: v[4] for k, v in lv.items()}
    check("V28 grid: the partials sit at least 6 per cent off on North Central; members uncalled converges",
          abs(nc["partial: class factor on every new NC premise"] - 210) / 210 >= 0.06
          and abs(nc["partial: uncalled draw on every new NC premise"] - 210) / 210 >= 0.06
          and lv["members uncalled in 2027 too"] == EXPECT["answer"], lv)
    # lenders' zone-share basis on ERCOT's outlook (the declared wrong-basis distractor)
    ws = load_workbook(os.path.join(P.d, "ercot_zone_peak_outlook_2027.xlsx"), read_only=True)["Weather zones"]
    out = {r[0]: r for r in list(ws.iter_rows(values_only=True))[4:] if r and r[0] in BOOKS}
    e26 = P.st[[d.year == 2026 for d in P.st.index.get_level_values(0)]][BOOKS].sum()
    lend = {b: e26[b] / (out[b][2] * 1000.0) * out[b][4] - hedges[b] for b in BOOKS}
    la = vec(level(lend))
    check("V29 the lenders' basis is Coast-led, puts North Central under 150, and is neither the answer nor rung 0",
          max(range(8), key=lambda k: la[k]) == 0 and la[4] < 150 and la != EXPECT["answer"] and la != EXPECT["R0"], la)
    # clean-data test: the enrolment repaired to one row per premise
    rows_rep = book27.groupby("book")["md"].sum().reindex(BOOKS).to_dict()
    r1_rep = vec(level(E(rows_rep, 0.0)))
    check("V30 clean-data test: with the enrolment repaired, rung 1 becomes rung 2; rung 3 and the answer stand",
          r1_rep == EXPECT["R2"] and vec(M["L"]["R3"]) == EXPECT["R3"] and vec(M["L"]["R4"]) == EXPECT["answer"])
    return {"grid": lv, "lenders": la, "new_general_nc_kw": newgen}


def gates(d, meta_path):
    files = sorted(os.listdir(d))
    fmts = {os.path.splitext(f)[1] for f in files}
    import pyarrow.parquet as pq
    n = pq.ParquetFile(os.path.join(d, "idr_hourly_reads_summers_2017_2026.parquet")).metadata.num_rows
    check("V31 input gates: 10 or more files, 3 or more formats, a file of 25,000 or more rows",
          len(files) >= 10 and len(fmts) >= 3 and n >= 25000, (len(files), sorted(fmts), n))
    if meta_path:
        meta = json.load(open(meta_path))
        ds = meta["distractor_files"]
        check("V32 two or more distractors named in metadata.json, present, and never named in the pack",
              len(ds) >= 2 and all(x in files for x in ds)
              and not any(re.search(rb"(?i)distractor", open(os.path.join(d, f), "rb").read()) for f in files))
    em = [f for f in files if f.endswith((".csv", ".md")) and "—" in open(os.path.join(d, f), encoding="utf-8").read()]
    check("V33 no em dash in any text file", not em, em)


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    out_json = sys.argv[sys.argv.index("--json") + 1] if "--json" in sys.argv else None
    if out_json in args:
        args.remove(out_json)
    d = args[0]
    meta = args[1] if len(args) > 1 else None
    P = Pack(d)
    M = main_call(P)
    C = corpus(P, M)
    twins_and_credits(P, M)
    A = asks(d)
    G = grid_and_lenders(P, M)
    gates(d, meta)
    n_fail = sum(1 for _, ok, _ in RESULTS if not ok)
    print(f"\n{len(RESULTS)} checks, {n_fail} failed. Answer {vec(M['L']['R4'])}; post-block "
          f"{ {b: round(v, 1) for b, v in M['post'].items()} }")
    if out_json:
        json.dump({"checks": RESULTS, "answer": vec(M["L"]["R4"]), "post": {b: round(v, 3) for b, v in M["post"].items()},
                   "loads": {b: round(M["E"]["R4"][b] + M["hedges"][b], 3) for b in BOOKS},
                   "exposure": {b: round(M["E"]["R4"][b], 3) for b in BOOKS}, "u": M["u"], "cf": M["cf"],
                   "corridor": M["corridor"], "estimators": M["ests"], "rivals": C["rivals"], "asks": A, "grid": G},
                  open(out_json, "w"), indent=1, default=str)
    sys.exit(1 if n_fail else 0)


if __name__ == "__main__":
    main()
