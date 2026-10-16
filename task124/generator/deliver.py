"""The golden deliverables, computed from the shipped pack only: lslr_2027_programme.xlsx (the zone schedule, the
unknown-line shares, the 2026 contract payments and the access agreements) and lslr_2027_board_paper.pdf (the call,
the schedule chart and the same tables), plus generator/figures.json for the write-up."""
from __future__ import annotations

import io
import json
import math
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import analysis as AN
import common as K
import load as L
import postprocess as PP

QUARTERS = (1, 2, 3, 4)
ERA_LABEL = ("Built before 1930", "Built 1930 to 1949", "Built 1950 to 1988")


def round_to(x, step):
    return int(math.floor(x / step + 0.5) * step)


def records_shares(inv):
    """Share of the tap card lines built before 1989 whose card reads LEAD or GALV, by era."""
    t = inv[(inv.record != "") & (inv.year_built > 0) & (inv.year_built < K.AGE_RULE_YEAR)]
    e = t.year_built.map(K.era_of)
    return {k: float(t[e == k].record.isin(["LEAD", "GALV"]).mean()) for k in range(3)}


def compute(target):
    inv, ver = L.inventory(target), L.verifications(target)
    prices, mob = L.unit_costs(target)
    ap = L.appropriation(target)
    s = AN.sample(inv, ver)
    cls = AN.street_class(inv)
    s["sc"] = s.segment.map(cls)
    sh = AN.kind_era_shares(s, cls)
    ub = AN.unknown_by_class(inv, cls)
    pu = AN.class_era_shares(inv, s, cls)
    sig = AN.share_sigma(s, cls, inv)
    rem = AN.remaining(inv)
    z = AN.zone_costs(rem, pu, ap, prices, mob)
    whole = s.coverage >= 0.999
    p0 = AN.shares(s)
    z0 = AN.zone_costs(rem, p0, ap, prices, mob)
    pw = AN.shares(s, whole)
    zw = AN.zone_costs(rem, pw, ap, prices, mob)
    prec = records_shares(inv)
    zr = AN.zone_costs(rem, prec, ap, prices, mob)
    pp = AN.class_era_shares(inv, s, cls, pooled=True)
    zp = AN.zone_costs(rem, pp, ap, prices, mob)
    # the evidence: verified no-tap-card lines by the class of their main, by era and by route
    ev = {}
    for k in ("LG", "C"):
        for e in range(3):
            for name, m in (("whole", whole), ("other", ~whole)):
                t = s[(s.sc == k) & (s.era == e) & m]
                ev[(k, e, name)] = (len(t), int(t.hit.sum()))
    unk = inv[inv.cls == "LEAD STATUS UNKNOWN"]
    unk_e = unk.year_built.map(K.era_of)
    car = inv[(inv.record != "") & (inv.year_built > 0) & (inv.year_built < K.AGE_RULE_YEAR)]
    car_e = car.year_built.map(K.era_of)
    mix = {"unknown": {e: float((unk[unk_e == e].segment.map(cls) == "LG").mean()) for e in range(3)},
           "verified": {e: float((s[s.era == e].sc == "LG").mean()) for e in range(3)},
           "whole": {e: float((s[whole & (s.era == e)].sc == "LG").mean()) for e in range(3)},
           "carded": {e: float((car[car_e == e].segment.map(cls) == "LG").mean()) for e in range(3)}}
    # ask A
    c = L.completions(target)
    co = L.change_orders(target)
    latest = c.sort_values(["site_id", "closed_on", "row_status"]).drop_duplicates("site_id", keep="last")
    latest = latest.assign(q=latest.local_date.map(K.quarter_of),
                           paid=latest.paid_base + latest.paid_extra_ft + latest.paid_restoration)
    site_q = dict(zip(latest.site_id, zip(latest.contract_no, latest.q)))
    co_q = {}
    for no, sites, amount in zip(co["CO no."], co["sites"], co["Amount ($)"]):
        keys = {site_q[sid] for sid in sites}
        assert len(keys) == 1, no
        co_q.setdefault(keys.pop(), []).append(float(amount))
    pay, n_done = {}, {}
    for con in (x[0] for x in K.CONTRACTS_2026):
        for qq in QUARTERS:
            m = (latest.contract_no == con) & (latest.q == qq)
            n = int(m.sum())
            pay[(con, qq)] = (float(latest.paid[m].sum()) + sum(co_q.get((con, qq), []))) / n
            n_done[(con, qq)] = n
    # ask B: a line has a signed agreement when the agreement system holds a signed form for it, or the portal shows
    # the owner's e-signed form and the agreement system holds no form for the line at all
    ag = L.agreements(target)
    tr = L.outreach(target)
    zone_of = dict(zip(inv.sl_id, inv.zone))
    signed = set(ag[ag.status == "SIGNED"].line)
    held = set(ag.line)
    portal_only = set(tr[tr.portal_status == "E-SIGNED"].sl_id) - held
    have = signed | portal_only
    agree = {zz: sum(1 for x in have if zone_of[x] == zz) for zz in K.PLAN_ORDER}
    cov = AN.coverage(inv, ver)
    return dict(s=s, whole=whole, cls=cls, sh=sh, ub=ub, pu=pu, sig=sig, z=z, p0=p0, z0=z0, pw=pw, zw=zw, prec=prec,
                zr=zr, pp=pp, zp=zp, ev=ev, mix=mix, ap=ap, pay=pay, n_done=n_done, agree=agree,
                portal_only=len(portal_only), n_segments=int((cov >= 0.999).sum()), n_sample=len(s), mob=mob,
                prices=prices, rem=rem)


def _pct(x):
    return "%d" % int(math.floor(100 * x + 0.5))


def figures(F):
    z = F["z"]
    last = z.iloc[-1]
    contractor = {x[0]: x[1] for x in K.CONTRACTS_2026}
    ub, sh, ev = F["ub"], F["sh"], F["ev"]
    total = float(z.cum_cost.iloc[-1])

    def cell(k, e):
        n = ev[(k, e, "whole")][0] + ev[(k, e, "other")][0]
        h = ev[(k, e, "whole")][1] + ev[(k, e, "other")][1]
        return {"verified": "{:,}".format(n), "found_lead_or_grr": "{:,}".format(h), "per_cent": "%.1f" % (100 * h / n)}
    out = {
        "call": {"last_zone": last.zone, "last_zone_name": K.ZONE_NAME[last.zone],
                 "last_finish_year": int(last.finish_year), "meets_deadline": bool(z.finish_year.max() <= 2037),
                 "deadline": "1 November 2037"},
        "appropriation_million": "%.1f" % (F["ap"] / 1e6),
        "total_plan_cost_million": "%d" % round_to(total / 1e6, 1),
        "programme_years": "%.1f" % (math.floor(10 * total / F["ap"] + 0.5) / 10),
        "era_shares_lead_status_unknown": {ERA_LABEL[e]: "%.2f" % (math.floor(100 * F["pu"][e] + 0.5) / 100)
                                           for e in range(3)},
        "kind_era_shares": {"lead_or_galvanized_cards": {ERA_LABEL[e]: cell("LG", e) for e in range(3)},
                            "copper_or_plastic_cards": {ERA_LABEL[e]: cell("C", e) for e in range(3)}},
        "unknown_by_class": {ERA_LABEL[e]: {"on_lead_or_galvanized_mains": "{:,}".format(ub[e]["LG"]),
                                            "on_copper_or_plastic_mains": "{:,}".format(ub[e]["C"]),
                                            "lead_status_unknown": "{:,}".format(ub[e]["LG"] + ub[e]["C"])}
                             for e in range(3)},
        "share_on_lead_or_galvanized_mains_pct": {k: {ERA_LABEL[e]: _pct(F["mix"][k][e]) for e in range(3)}
                                                  for k in ("unknown", "verified", "whole", "carded")},
        "verified_no_tap_card_lines": "{:,}".format(F["n_sample"]),
        "segments_verified_whole": F["n_segments"],
        "mobilisation": "{:,}".format(int(F["mob"])),
        "zones": {},
        "default_path": {"era_shares_every_verification": {ERA_LABEL[e]: "%.3f" % F["p0"][e] for e in range(3)},
                         "last_finish_year": int(F["z0"].finish_year.max())},
        "whole_segment_path": {"era_shares": {ERA_LABEL[e]: "%.3f" % F["pw"][e] for e in range(3)},
                               "last_finish_year": int(F["zw"].finish_year.max())},
        "tap_card_path": {"era_shares": {ERA_LABEL[e]: "%.3f" % F["prec"][e] for e in range(3)},
                          "last_finish_year": int(F["zr"].finish_year.max())},
        "pooled_era_path": {"share_lead_or_galvanized_mains_all_eras": "%.3f" % AN.class_shares(F["s"], F["cls"])["LG"],
                            "zones_whose_lines_differ": sum(
                                1 for a, b in zip(F["zp"].replacements, z.replacements)
                                if round_to(a, 10) != round_to(b, 10)),
                            "last_finish_year": int(F["zp"].finish_year.max())},
        "payments_2026_per_full_replacement": {},
        "full_replacements_2026": {},
        "agreements_signed": {zz: F["agree"][zz] for zz in K.PLAN_ORDER},
    }
    prev = 0.0
    for i, r in enumerate(z.itertuples(), 1):
        out["zones"][r.zone] = {"priority": i, "name": K.ZONE_NAME[r.zone],
                                "lead_and_grr_to_replace": "{:,}".format(round_to(r.replacements, 10)),
                                "plan_cost": "{:,}".format(round_to(r.cost, 100_000)),
                                "plan_cost_million": "%.1f" % (round_to(r.cost, 100_000) / 1e6),
                                "start_year": K.FIRST_YEAR + int(math.floor(prev / F["ap"])),
                                "finish_year": int(r.finish_year)}
        prev = r.cum_cost
    for con in contractor:
        out["payments_2026_per_full_replacement"][con] = {
            "contractor": contractor[con],
            **{"Q%d" % qq: "{:,}".format(int(math.floor(F["pay"][(con, qq)] + 0.5))) for qq in QUARTERS}}
        out["full_replacements_2026"][con] = {"Q%d" % qq: F["n_done"][(con, qq)] for qq in QUARTERS}
    return out


# ------------------------------------------------------------------------------------------------ workbook
HEAD_FILL = PatternFill("solid", fgColor="1F4E5A")
SUB_FILL = PatternFill("solid", fgColor="DCE6EA")


def _head(ws, row, cols, widths=None):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 32
    for j, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def write_xlsx(path, F, fig):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Zone schedule"
    ws["A1"] = ("Every zone finished by %d at $%s million a year: lead service line replacement schedule, 2027 to 2037"
                % (fig["call"]["last_finish_year"], fig["appropriation_million"]))
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("%s, LSLR programme office. Plan revision 2 zones in priority order; plan cost in 2027 dollars at the "
                "2027 unit cost schedule. Prepared 11 January 2027 for the board meeting of 19 January 2027."
                % K.DISTRICT)
    ws["A2"].font = Font(italic=True, size=9)
    _head(ws, 4, ["Priority", "Zone", "Zone name", "Lead and GRR lines to replace", "Plan cost ($)",
                  "Work starts (programme year)", "Finish year"], [8, 8, 18, 14, 14, 13, 9])
    r = 5
    for zz, d in fig["zones"].items():
        vals = [d["priority"], zz, d["name"], int(d["lead_and_grr_to_replace"].replace(",", "")),
                int(d["plan_cost"].replace(",", "")), d["start_year"], d["finish_year"]]
        fmts = [None, None, None, "#,##0", "$#,##0", "0", "0"]
        for j, (v, f) in enumerate(zip(vals, fmts), 1):
            cell = ws.cell(row=r, column=j, value=v)
            if f:
                cell.number_format = f
        if zz == fig["call"]["last_zone"]:
            for j in range(1, 8):
                ws.cell(row=r, column=j).font = Font(bold=True)
        r += 1
    ws.cell(row=r, column=3, value="All twelve zones").font = Font(bold=True)
    ws.cell(row=r, column=5, value="$%s million" % fig["total_plan_cost_million"]).alignment = Alignment(
        horizontal="right")
    ws.cell(row=r, column=6, value="%s years" % fig["programme_years"]).alignment = Alignment(horizontal="right")
    for j in range(1, 8):
        ws.cell(row=r, column=j).fill = SUB_FILL
    r += 2
    notes = [
        "Last zone: %s %s, finished in programme year %d, inside the %s deadline. The twelve zones cost $%s million, "
        "%s years of the $%s million appropriation." % (
            fig["call"]["last_zone"], fig["call"]["last_zone_name"], fig["call"]["last_finish_year"],
            fig["call"]["deadline"], fig["total_plan_cost_million"], fig["programme_years"],
            fig["appropriation_million"]),
        "Lines to replace: the zone's Lead and GRR lines plus its Lead Status Unknown lines at the share of lead or GRR "
        "expected among the district's Lead Status Unknown lines of their construction era (plan 4.3; sheet "
        "Unknown-line shares).",
        "A zone is finished in the programme year in which the running total of the appropriation covers the running "
        "total of plan cost through it (plan 5.3); its work starts when the zone before it is finished (plan 5.2). "
        "Lines rounded to the nearest ten and costs to the nearest $100,000, from unrounded expected lines.",
    ]
    for t in notes:
        ws.cell(row=r, column=1, value=t).font = Font(size=9)
        r += 1
    ws.freeze_panes = "A5"

    ws2 = wb.create_sheet("Unknown-line shares")
    ws2["A1"] = "Share of the district's Lead Status Unknown lines expected to prove lead or GRR, by construction era"
    ws2["A1"].font = Font(bold=True, size=12)
    ws2["A2"] = ("All fourteen zones. Each unknown line is costed at the rate found on the verified lines with no tap card "
                 "on mains of the same kind built in the same era: mains whose tap cards read lead or galvanized, and "
                 "mains whose cards read only copper or plastic.")
    ws2["A2"].font = Font(italic=True, size=9)
    _head(ws2, 4, ["Construction era", "Unknown lines on lead or galvanized mains", "Unknown lines on copper or "
                   "plastic mains", "Lead Status Unknown lines", "Share among Lead Status Unknown lines"],
          [20, 15, 15, 14, 14, 14, 11])
    for e in range(3):
        u = fig["unknown_by_class"][ERA_LABEL[e]]
        row = [ERA_LABEL[e], int(u["on_lead_or_galvanized_mains"].replace(",", "")),
               int(u["on_copper_or_plastic_mains"].replace(",", "")), int(u["lead_status_unknown"].replace(",", "")),
               float(fig["era_shares_lead_status_unknown"][ERA_LABEL[e]])]
        for j, v in enumerate(row, 1):
            cell = ws2.cell(row=5 + e, column=j, value=v)
            cell.number_format = "0.00" if j == 5 else ("#,##0" if j > 1 else "General")
    ws2.cell(row=9, column=1, value="Verified lines with no tap card, first result").font = Font(bold=True, size=9)
    _head(ws2, 10, ["Main's tap cards", "Construction era", "On the replaced mains: verified", "found lead or GRR",
                    "Other verifications: verified", "found lead or GRR", "Found lead or GRR (%)"])
    r = 11
    for k, label in (("LG", "Lead or galvanized"), ("C", "Copper or plastic only")):
        for e in range(3):
            nw, hw = F["ev"][(k, e, "whole")]
            no, ho = F["ev"][(k, e, "other")]
            row = [label, ERA_LABEL[e], nw, hw, no, ho, round(100 * (hw + ho) / (nw + no), 1)]
            for j, v in enumerate(row, 1):
                cell = ws2.cell(row=r, column=j, value=v)
                if j == 7:
                    cell.number_format = "0.0"
                elif j > 2:
                    cell.number_format = "#,##0"
            r += 1
    r += 1
    ws2.cell(row=r, column=1, value="Lines on lead or galvanized mains, per cent of each group").font = Font(bold=True,
                                                                                                          size=9)
    _head(ws2, r + 1, ["Construction era", "Lead Status Unknown lines", "Verified lines with no tap card",
                       "Lines on the replaced mains", "Tap card lines"])
    for e in range(3):
        mx = fig["share_on_lead_or_galvanized_mains_pct"]
        for j, v in enumerate([ERA_LABEL[e], int(mx["unknown"][ERA_LABEL[e]]), int(mx["verified"][ERA_LABEL[e]]),
                               int(mx["whole"][ERA_LABEL[e]]), int(mx["carded"][ERA_LABEL[e]])], 1):
            ws2.cell(row=r + 2 + e, column=j, value=v).number_format = "0" if j > 1 else "General"
    ws2.cell(row=r + 6, column=1, value=(
        "On lead or galvanized mains the share found falls with the era the line was built, so each era is costed at "
        "its own rate; one rate for each kind of main across all eras (%s on lead or galvanized mains) would change the "
        "rounded lines in %d of the twelve zones." % (
            fig["pooled_era_path"]["share_lead_or_galvanized_mains_all_eras"],
            fig["pooled_era_path"]["zones_whose_lines_differ"]))).font = Font(size=9)
    ws2.cell(row=r + 7, column=1, value=(
        "The lines verified so far and the surviving tap cards sit on lead or galvanized mains far more often than the "
        "lines still unknown, so a rate taken from any of them across a whole era misstates the unknown lines: at the "
        "shares across every verification the plan would run to %d, at the shares on the replaced mains to %d, "
        "and at the tap card lines' share to %d." % (
            fig["default_path"]["last_finish_year"], fig["whole_segment_path"]["last_finish_year"],
            fig["tap_card_path"]["last_finish_year"]))).font = Font(size=9)
    ws2.column_dimensions["G"].width = 11
    ws2.freeze_panes = "A5"

    ws3 = wb.create_sheet("2026 contract payments")
    ws3["A1"] = "Paid per full replacement completed, 2026 replacement contracts"
    ws3["A1"].font = Font(bold=True, size=12)
    ws3["A2"] = ("Bid price, extra service length and restoration on each site's latest closing, plus each approved "
                 "change order once, over the full replacements completed in the quarter (Eastern time). Dollars.")
    ws3["A2"].font = Font(italic=True, size=9)
    _head(ws3, 4, ["Contract", "Contractor", "Q1 2026", "Q2 2026", "Q3 2026", "Q4 2026"], [13, 26, 11, 11, 11, 11])
    r = 5
    for con, d in fig["payments_2026_per_full_replacement"].items():
        ws3.cell(row=r, column=1, value=con)
        ws3.cell(row=r, column=2, value=d["contractor"])
        for j, q in enumerate(QUARTERS, 3):
            ws3.cell(row=r, column=j, value=int(d["Q%d" % q].replace(",", ""))).number_format = "#,##0"
        r += 1
    r += 1
    ws3.cell(row=r, column=1, value="Full replacements completed").font = Font(bold=True, size=9)
    r += 1
    for con, d in fig["full_replacements_2026"].items():
        ws3.cell(row=r, column=1, value=con)
        ws3.cell(row=r, column=2, value=fig["payments_2026_per_full_replacement"][con]["contractor"])
        for j, q in enumerate(QUARTERS, 3):
            ws3.cell(row=r, column=j, value=d["Q%d" % q]).number_format = "#,##0"
        r += 1
    ws3.cell(row=r + 1, column=1, value=(
        "Tri-County Underground was paid under vendor V-20877 to 12 July 2026 and V-31560 from 13 July 2026; both are "
        "contract 2026-LSL-02.")).font = Font(size=9)
    ws3.freeze_panes = "C5"

    ws4 = wb.create_sheet("Access agreements")
    ws4["A1"] = "Service lines with a signed access agreement, by zone, 8 January 2027"
    ws4["A1"].font = Font(bold=True, size=12)
    ws4["A2"] = ("Distinct service lines open in the plan with a right-of-entry agreement signed and not withdrawn: "
                 "forms in the agreement system, and forms the owner e-signed on the portal that the agreement system "
                 "does not hold; counted in the zone the inventory gives the line today.")
    ws4["A2"].font = Font(italic=True, size=9)
    _head(ws4, 4, ["Priority", "Zone", "Zone name", "Service lines with a signed agreement"], [8, 8, 18, 16])
    for i, zz in enumerate(K.PLAN_ORDER, 1):
        for j, v in enumerate([i, zz, K.ZONE_NAME[zz], fig["agreements_signed"][zz]], 1):
            cell = ws4.cell(row=4 + i, column=j, value=v)
            if j == 4:
                cell.number_format = "#,##0"
    ws4.freeze_panes = "A5"
    for sheet, area, titles in ((ws, "A1:G%d" % ws.max_row, "4:4"), (ws2, "A1:G%d" % ws2.max_row, None),
                                (ws3, "A1:F%d" % ws3.max_row, None), (ws4, "A1:D%d" % ws4.max_row, "4:4")):
        sheet.print_area = area
        if titles:
            sheet.print_title_rows = titles
        sheet.page_setup.orientation = "landscape" if sheet in (ws, ws2) else "portrait"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)
    PP.normalize_ooxml(path, K.PEOPLE["programme"], K.PEOPLE["programme"], "2027-01-08T15:02:11Z",
                       "2027-01-11T22:18:40Z", excel=True)


# ----------------------------------------------------------------------------------------------- the chart
def chart_png(F, fig):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    for f in ("LiberationSans-Regular.ttf", "LiberationSans-Bold.ttf"):
        font_manager.fontManager.addfont("/usr/share/fonts/truetype/liberation/" + f)
    plt.rcParams.update({"font.family": "Liberation Sans", "font.size": 8.5, "svg.hashsalt": "task124"})
    z = F["z"]
    ap = F["ap"]
    fig_, ax = plt.subplots(figsize=(7.0, 3.9), dpi=200)
    n = len(z)
    prev = 0.0
    for i, r in enumerate(z.itertuples()):
        y = n - 1 - i
        x0, x1 = K.FIRST_YEAR + prev / ap, K.FIRST_YEAR + r.cum_cost / ap
        last = i == n - 1
        ax.barh(y, x1 - x0, left=x0, height=0.62, color="#1F4E5A" if not last else "#C0562B", edgecolor="none")
        near = x1 > K.LAST_YEAR + 1 - 0.65      # a label past the bar would sit on the deadline line
        ax.text(x0 - 0.08 if near else x1 + 0.08, y, str(int(r.finish_year)), va="center",
                ha="right" if near else "left", fontsize=7.8, fontweight="bold" if last else "normal",
                color="#1d2329")
        prev = r.cum_cost
    ax.set_yticks(range(n))
    ax.set_yticklabels(["%s %s" % (r.zone, K.ZONE_NAME[r.zone]) for r in z.itertuples()][::-1])
    dl = K.LAST_YEAR + 1
    ax.axvline(dl, color="#8a1c1c", linestyle="--", linewidth=1.1)
    ax.text(dl - 0.06, n - 0.35, "Deadline\n1 Nov 2037", ha="right", va="top", fontsize=7.6, color="#8a1c1c")
    ax.set_xlim(K.FIRST_YEAR, K.LAST_YEAR + 2)
    ax.set_xticks([y + 0.5 for y in range(K.FIRST_YEAR, K.LAST_YEAR + 2)])
    ax.set_xticklabels([str(y) for y in range(K.FIRST_YEAR, K.LAST_YEAR + 2)])
    for y in range(K.FIRST_YEAR, K.LAST_YEAR + 3):
        ax.axvline(y, color="#e3e7ea", linewidth=0.6, zorder=0)
    ax.tick_params(axis="x", length=0)
    ax.set_xlabel("Programme year (each season 1 April to 31 October), at $%s million a year"
                  % fig["appropriation_million"])
    ax.set_title("Every zone finished by %d: %s %s last, inside the deadline"
                 % (fig["call"]["last_finish_year"], fig["call"]["last_zone"], fig["call"]["last_zone_name"]),
                 fontsize=9.5, loc="left", fontweight="bold")
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    fig_.tight_layout()
    buf = io.BytesIO()
    fig_.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig_)
    return buf.getvalue()


# ----------------------------------------------------------------------------------------------- board paper
def write_pdf(path, F, fig):
    from reportlab import rl_config
    rl_config.invariant = 1
    from reportlab.lib.units import inch
    from reportlab.platypus import Image, KeepTogether, Paragraph, Spacer
    import docs as D
    doc = D._doc(path, "Lead service line replacement schedule, 2027 to 2037", K.PEOPLE["programme"])
    z = F["z"]
    st = [Paragraph("Every replacement zone finished by %d at $%s million a year" % (
        fig["call"]["last_finish_year"], fig["appropriation_million"]), D.H1),
          Paragraph("Board of Directors, 19 January 2027, item 6: capital improvement plan 2027 to 2037. From %s, LSLR "
                    "Programme Manager, with %s, District Engineer. Prepared 11 January 2027." % (
                        K.PEOPLE["programme"], K.PEOPLE["engineer"]), D.META)]
    st.append(Paragraph(
        "At the capital plan's $%s million a year from 2027, plan revision 2 finishes all twelve zones inside the "
        "Lead and Copper Rule Improvements deadline. The last zone, %s %s, is finished in programme year %d, before 1 "
        "November 2037: the plan cost of the twelve zones is $%s million, %s years of the appropriation. The board can "
        "adopt the construction line as drafted." % (fig["appropriation_million"], fig["call"]["last_zone"],
                                                      fig["call"]["last_zone_name"], fig["call"]["last_finish_year"],
                                                      fig["total_plan_cost_million"], fig["programme_years"]), D.BODY))
    img = chart_png(F, fig)
    st.append(Image(io.BytesIO(img), width=6.7 * inch, height=6.7 * inch * 3.9 / 7.0))
    st.append(Paragraph("Zones in priority order; each bar runs from the zone's start to its finish on the running total "
                        "of plan cost against the appropriation (plan 5.2 and 5.3).", D.NOTE))
    rows = [["Priority", "Zone", "Lead and GRR lines to replace", "Plan cost ($ million)", "Finish year"]]
    for zz, d in fig["zones"].items():
        rows.append([str(d["priority"]), "%s %s" % (zz, d["name"]), d["lead_and_grr_to_replace"],
                     d["plan_cost_million"], str(d["finish_year"])])
    st.append(KeepTogether([Paragraph("Zone schedule", D.H2), D._table(rows, [0.7 * inch, 1.9 * inch, 1.5 * inch,
                                                                              1.3 * inch, 0.9 * inch],
                                                                       right=(2, 3, 4))]))
    st.append(Paragraph("Lines to the nearest ten and costs to the nearest $100,000, 2027 dollars at the 2027 unit cost "
                        "schedule with one $%s mobilisation per zone." % fig["mobilisation"], D.NOTE))
    unk = fig["era_shares_lead_status_unknown"]
    ubc = fig["unknown_by_class"]
    mx = fig["share_on_lead_or_galvanized_mains_pct"]
    dflt = fig["default_path"]["era_shares_every_verification"]
    kes = fig["kind_era_shares"]
    lg, cu = kes["lead_or_galvanized_cards"], kes["copper_or_plastic_cards"]
    E0, E1, E2 = ERA_LABEL

    def none_or(c):
        return "none of %s" % c["verified"] if c["found_lead_or_grr"] == "0" else "%s of %s" % (
            c["found_lead_or_grr"], c["verified"])
    st.append(Paragraph("How the unknown lines are costed", D.H2))
    st.append(Paragraph(
        "Plan 4.3 carries each Lead Status Unknown line at the share of lead or GRR expected among the district's Lead "
        "Status Unknown lines of its construction era. Services were laid with the main they tap, so the tap cards "
        "that survive on a main show what kind of street a line without a card sits on, and who laid the services "
        "changed with the decades. Among the lines with no tap card that crews have verified, those on mains whose "
        "cards read lead or galvanized proved lead or GRR at %s per cent when built before 1930, %s per cent when built "
        "1930 to 1949 and %s per cent when built 1950 to 1988; on mains whose cards read only copper or plastic, %s, "
        "%s and %s. The rates are alike on the replaced mains and on single-line visits. The lines still unknown sit "
        "on lead or galvanized mains far less often than the lines verified so far (%s, %s and %s per cent of the "
        "unknown lines by era, against %s, %s and %s per cent of the verified ones), so the %s, %s and %s unknown "
        "lines are expected to prove lead or GRR at shares of %s, %s and %s. Priced at the shares found across every "
        "verification on file (%s, %s and %s) the same plan would run to %d; at the shares on the replaced mains, to "
        "%d; at the share of tap cards that read lead or galvanized, to %d." % (
            lg[E0]["per_cent"], lg[E1]["per_cent"], lg[E2]["per_cent"], none_or(cu[E0]), none_or(cu[E1]),
            none_or(cu[E2]), mx["unknown"][E0], mx["unknown"][E1], mx["unknown"][E2], mx["verified"][E0],
            mx["verified"][E1], mx["verified"][E2], ubc[E0]["lead_status_unknown"], ubc[E1]["lead_status_unknown"],
            ubc[E2]["lead_status_unknown"], unk[E0], unk[E1], unk[E2], dflt[E0], dflt[E1], dflt[E2],
            fig["default_path"]["last_finish_year"], fig["whole_segment_path"]["last_finish_year"],
            fig["tap_card_path"]["last_finish_year"]), D.BODY))
    rows = [["Contract", "Contractor", "Q1", "Q2", "Q3", "Q4"]]
    for con, d in fig["payments_2026_per_full_replacement"].items():
        rows.append([con, d["contractor"]] + [d["Q%d" % q] for q in QUARTERS])
    st.append(KeepTogether([
        Paragraph("What a full replacement cost in 2026", D.H2),
        Paragraph("Paid per full replacement completed, by quarter of 2026 (Eastern time), in dollars: bid price, extra "
                  "service length and restoration on each site's latest closing, with each approved change order "
                  "counted once.", D.BODY),
        D._table(rows, [1.0 * inch, 2.1 * inch, 0.8 * inch, 0.8 * inch, 0.8 * inch, 0.8 * inch],
                 right=(2, 3, 4, 5))]))
    st.append(Paragraph("Tri-County Underground's vendor number changed on 13 July 2026 after its acquisition; its "
                        "contract did not, and both vendor numbers are counted under 2026-LSL-02.", D.NOTE))
    rows = [["Priority", "Zone", "Service lines with a signed access agreement"]]
    for i, zz in enumerate(K.PLAN_ORDER, 1):
        rows.append([str(i), "%s %s" % (zz, K.ZONE_NAME[zz]), "{:,}".format(fig["agreements_signed"][zz])])
    st.append(KeepTogether([
        Paragraph("Access agreements in hand", D.H2),
        D._table(rows, [0.7 * inch, 1.9 * inch, 2.0 * inch], right=(2,))]))
    st.append(Paragraph("Distinct service lines open in the plan with a right-of-entry agreement signed and not "
                        "withdrawn at 8 January 2027, including forms owners e-signed on the portal that the agreement "
                        "system does not hold, in the zone the inventory gives each line under plan revision 2; some "
                        "forms signed before the revision print a different zone.", D.NOTE))
    foot = D._frame("%s  |  Board of Directors, 19 January 2027" % K.DISTRICT, "LSLR schedule 2027 to 2037")
    doc.build(st, onFirstPage=foot, onLaterPages=foot)
    PP.normalize_pdf(path, b"D:20270111161204-05'00'", b"D:20270111173518-05'00'")


def main(target, golden, figures_path, world=None):
    os.makedirs(golden, exist_ok=True)
    F = compute(target)
    fig = figures(F)
    write_xlsx(os.path.join(golden, K.DELIVERABLES[0]), F, fig)
    write_pdf(os.path.join(golden, K.DELIVERABLES[1]), F, fig)
    with open(figures_path, "w") as fh:
        json.dump(fig, fh, indent=1)
        fh.write("\n")
    return F, fig
