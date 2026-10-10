"""task124 pack build.

    python3 task124/generator/build_pack.py --out <dir for target files> --meta <path of metadata.json>
                                            [--record <path for the build record json>]

Seeded and deterministic: two runs write byte-identical files. Every assertion in checks.py runs before a file is
written; the pack gates run after."""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime
from zoneinfo import ZoneInfo

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, HERE)

import checks as C  # noqa: E402
import docs2 as D2  # noqa: E402
import documents as D  # noqa: E402
import pipeline  # noqa: E402
import writers as Wr  # noqa: E402
from analysis import Analysis  # noqa: E402
from common import SEED  # noqa: E402

SCRUB = os.path.join(REPO, ".claude", "skills", "reduce-house-fixes", "scripts", "scrub_producer_metadata.py")
CT = ZoneInfo("America/Chicago")
FLOOR, CEILING = "2024-01-01", "2027-04-12"

# in-fiction write time of every shipped file, local Central time
WRITTEN = {
    Wr.SPINE: (2027, 4, 9, 18, 52), Wr.SETTLED: (2027, 4, 9, 18, 31), Wr.PEAKLIST: (2026, 10, 2, 9, 14),
    Wr.ENROL: (2027, 4, 9, 19, 3), Wr.PREMISES: (2027, 4, 9, 19, 11), Wr.SWITCHES: (2027, 4, 9, 19, 16),
    Wr.POSITIONS: (2027, 4, 9, 18, 18), Wr.CREDITS: (2027, 4, 9, 17, 48), Wr.ACCOUNTS: (2027, 4, 9, 17, 52),
    Wr.ROSTER: (2027, 3, 31, 16, 46), Wr.TERMS: (2025, 3, 3, 10, 12), Wr.POLICY: (2027, 2, 19, 14, 41),
    Wr.MINUTE: (2027, 3, 22, 16, 5), Wr.REPORT: (2027, 2, 2, 15, 38), Wr.QUOTES: (2027, 4, 9, 17, 5),
    Wr.PECOS: (2027, 4, 9, 15, 2), Wr.DECISIONS: (2027, 4, 9, 17, 21), Wr.CPTYS: (2027, 4, 8, 11, 37),
    Wr.BOOKMAP: (2027, 1, 4, 10, 2), Wr.BLOTTER: (2027, 4, 9, 18, 40), Wr.MATCHLOG: (2027, 4, 9, 18, 41),
    Wr.PORTFOLIOS: (2026, 11, 12, 13, 26), Wr.CALENDAR: (2026, 12, 1, 9, 48), Wr.PROCEDURES: (2027, 1, 15, 16, 7),
    Wr.TEMPS: (2026, 10, 5, 8, 55), Wr.DAYAHEAD: (2026, 10, 6, 14, 12), Wr.OUTLOOK: (2026, 12, 18, 10, 30), Wr.NOTES: (2027, 4, 12, 11, 20),
    Wr.LOG: (2027, 4, 12, 16, 45),
}
PRODUCERS = {Wr.TERMS: ("Sabine Crest Energy", "2025-03-03"), Wr.MINUTE: ("Sabine Crest Energy", "2027-03-22"),
             Wr.REPORT: ("Sabine Crest Energy", "2027-02-02"), Wr.POLICY: ("Tammy Ochoa", "2027-02-19"),
             Wr.PROCEDURES: ("Gregory Sheppard", "2027-01-15"), Wr.POSITIONS: ("Trade Operations", "2027-04-09"),
             Wr.ROSTER: ("Donald Lee", "2027-03-31"), Wr.PECOS: ("Pecos Power Brokerage", "2027-04-09"),
             Wr.OUTLOOK: ("ERCOT", "2026-12-18")}

SOURCES = [
    (Wr.SPINE, "MDM interval export, IDR premises, weekdays June to September 2017 to 2026, hours ending 11 to 20, run 9 April 2027."),
    (Wr.SETTLED, "ERCOT settlement extracts for the book by weather zone at true-up, June to September 2017 to 2026."),
    (Wr.PEAKLIST, "ERCOT's published summer system peaks, 2017 to 2026, as kept by Load Planning."),
    (Wr.ENROL, "Enrolment system, every enrolment row as of 9 April 2027."),
    (Wr.PREMISES, "Premise register: ERCOT ESI ID attributes with the NAICS code from the service application, 9 April 2027."),
    (Wr.SWITCHES, "Completed switch and move-in transactions, 1 October 2026 to 9 April 2027."),
    (Wr.POSITIONS, "Trade Operations position report, close of 9 April 2027."),
    (Wr.CREDITS, "Billing system: Business Saver credit lines, 2017 to 2026 seasons."),
    (Wr.ACCOUNTS, "Billing system: IDR summary-billed accounts and their ESI IDs, 9 April 2027."),
    (Wr.ROSTER, "Program Desk roster for the 2027 season, renewals closed 31 March 2027."),
    (Wr.TERMS, "Business Saver program terms, 2025 edition."),
    (Wr.POLICY, "Summer Supply Risk Policy, adopted 19 February 2027."),
    (Wr.MINUTE, "Risk Committee minute of 19 March 2027."),
    (Wr.REPORT, "Load Planning, summer 2026 close-out, 2 February 2027."),
    (Wr.QUOTES, "Trading: quotes received on the desk's quote line, 1 to 9 April 2027."),
    (Wr.PECOS, "Pecos Power Brokerage's offer sheet to Sabine Crest, as received 9 April 2027."),
    (Wr.DECISIONS, "Trading: quote decisions log, 1 to 9 April 2027."),
    (Wr.CPTYS, "Trade Operations counterparty master, 8 April 2027."),
    (Wr.BOOKMAP, "Trading: book to ERCOT load zone map."),
    (Wr.BLOTTER, "Trade Operations blotter: ERCOT fixed-price strips held for the retail book, as of 9 April 2027."),
    (Wr.MATCHLOG, "Trade Operations confirmation matching log, as of 9 April 2027."),
    (Wr.PORTFOLIOS, "Trading: portfolio crosswalk."),
    (Wr.CALENDAR, "Trading calendar for 2027."),
    (Wr.PROCEDURES, "Trading desk procedures, Rev. 3."),
    (Wr.TEMPS, "Daily temperatures by weather zone, June to September 2017 to 2026, from the weather vendor feed."),
    (Wr.DAYAHEAD, "Load Planning's day-ahead forecasts of the book by weather zone, summer 2026, as issued at 10:00 the day before."),
    (Wr.OUTLOOK, "ERCOT summer 2027 outlook by weather zone, December 2026 release."),
    (Wr.NOTES, "Load Planning field notes for the extracts."),
]

MAIN_FILES = [Wr.SPINE, Wr.SETTLED, Wr.PEAKLIST, Wr.ENROL, Wr.PREMISES, Wr.SWITCHES, Wr.POSITIONS, Wr.CREDITS,
              Wr.ACCOUNTS, Wr.ROSTER, Wr.TERMS, Wr.POLICY, Wr.MINUTE, Wr.REPORT, Wr.NOTES, Wr.TEMPS]
DEVICE_FILES = [Wr.QUOTES, Wr.PECOS, Wr.DECISIONS, Wr.CPTYS, Wr.BOOKMAP, Wr.BLOTTER, Wr.MATCHLOG, Wr.PORTFOLIOS,
                Wr.CALENDAR, Wr.PROCEDURES]


def scrub(path, producer, stamp):
    with tempfile.TemporaryDirectory() as td:
        tmp = os.path.join(td, os.path.basename(path))
        shutil.copy2(path, tmp)
        prod = producer.replace("&", "&amp;") if path.endswith((".docx", ".xlsx")) else producer
        r = subprocess.run([sys.executable, SCRUB, td, "--apply", "--producer", prod, "--stamp", stamp,
                            "--floor", FLOOR, "--ceiling", CEILING], capture_output=True, text=True)
        a = subprocess.run([sys.executable, SCRUB, td, "--floor", FLOOR, "--ceiling", CEILING], capture_output=True,
                           text=True)
        C.ck(f"H1 container scrubbed: {os.path.basename(path)}", a.returncode == 0 and "clean" in a.stdout,
             r.stdout + r.stderr + a.stdout)
        shutil.copyfile(tmp, path)


def doc_texts(out):
    from docx import Document
    from openpyxl import load_workbook
    from pypdf import PdfReader
    texts = {}
    for f in sorted(os.listdir(out)):
        p = os.path.join(out, f)
        if f.endswith(".pdf"):
            texts[f] = "\n".join(pg.extract_text() for pg in PdfReader(p).pages)
        elif f.endswith(".docx"):
            d = Document(p)
            texts[f] = "\n".join([x.text for x in d.paragraphs] + [c.text for t in d.tables for r in t.rows for c in r.cells])
        elif f.endswith(".xlsx"):
            wb = load_workbook(p, read_only=True)
            texts[f] = "\n".join(" ".join(str(c) for c in row if c is not None)
                                 for ws in wb.worksheets for row in ws.iter_rows(values_only=True))
        elif f.endswith(".md"):
            texts[f] = open(p).read()
    return texts


def write_all(w, an, out):
    os.makedirs(out, exist_ok=True)
    for f in os.listdir(out):
        os.remove(os.path.join(out, f))
    j = lambda f: os.path.join(out, f)  # noqa: E731
    rng = np.random.default_rng(SEED + 99)
    info = {"spine_rows": Wr.write_spine(w, j(Wr.SPINE))}
    Wr.write_settled(w, j(Wr.SETTLED))
    Wr.write_peaks(j(Wr.PEAKLIST))
    Wr.write_enrolment(w, j(Wr.ENROL))
    Wr.write_premises(w, j(Wr.PREMISES), rng)
    Wr.write_switches(w, j(Wr.SWITCHES), rng)
    acc = Wr.account_frame(w, rng)
    Wr.write_accounts(acc, j(Wr.ACCOUNTS))
    Wr.write_credits(w, j(Wr.CREDITS))
    Wr.write_quotes(w, j(Wr.QUOTES))
    Wr.write_decisions(w, j(Wr.DECISIONS), rng)
    Wr.write_cptys(j(Wr.CPTYS))
    Wr.write_bookmap(j(Wr.BOOKMAP))
    Wr.write_blotter(w, j(Wr.BLOTTER))
    Wr.write_matchlog(w, j(Wr.MATCHLOG))
    Wr.write_portfolios(j(Wr.PORTFOLIOS))
    Wr.write_calendar(j(Wr.CALENDAR))
    Wr.write_temps(w, j(Wr.TEMPS))
    Wr.write_dayahead(w, j(Wr.DAYAHEAD), np.random.default_rng(SEED + 98))
    D.terms(j(Wr.TERMS))
    D.minute(j(Wr.MINUTE))
    D.report(j(Wr.REPORT), an)
    D2.policy(j(Wr.POLICY))
    D2.procedures(j(Wr.PROCEDURES))
    D2.positions(j(Wr.POSITIONS), w)
    D2.roster(j(Wr.ROSTER), w)
    D2.pecos(j(Wr.PECOS), w)
    D2.outlook(j(Wr.OUTLOOK))
    D2.field_notes(j(Wr.NOTES))
    D2.extract_log(j(Wr.LOG), SOURCES)
    for f, (producer, stamp) in PRODUCERS.items():
        scrub(j(f), producer, stamp)
    for f, t in WRITTEN.items():
        ts = datetime(*t, tzinfo=CT).timestamp()
        os.utime(j(f), (ts, ts))
    return info


RULES = {"amendment supersedes": r"supersedes it", "inclusive P90": r"linear interpolation between closest ranks",
         "in service from 1 June": r"in service from 1 June", "system peak hour": r"ERCOT's settled summer peak",
         "lot rule": r"largest remaining uncovered exposure", "lenders' basis": r"zone-share basis",
         "block cap": r"400 MW", "eligibility": r"full June through September of interval reads",
         "baseline": r"ten most recent business days without a call", "version of record": r"latest amendment matched",
         "MWh weighting": r"weights each trade by its MWh", "approval on the day": r"approved counterparty on the day",
         "lowest accepted quote": r"lowest accepted quote", "conversion": r"hours of the notional",
         "book map by month": r"the book map gives", "replay method": r"times the book's enrolled maximum demand"}


def container_stamps(out):
    bad = []
    for f in sorted(os.listdir(out)):
        p = os.path.join(out, f)
        if f.endswith(".pdf"):
            b = open(p, "rb").read()
            got = re.findall(rb"/(CreationDate|ModDate) \(D:(\d{8})(\d{6})([+-]\d\d)'(\d\d)'\)", b)
            if len(got) != 2 or any(t == b"000000" or o not in (b"-05", b"-06") for _, _, t, o, _ in got):
                bad.append(f)
        elif f.endswith((".docx", ".xlsx")):
            with zipfile.ZipFile(p) as z:
                core = z.read("docProps/core.xml").decode()
                app = z.read("docProps/app.xml").decode()
            got = re.findall(r"<dcterms:(created|modified)[^>]*>(\d{4}-\d\d-\d\d)T(\d\d:\d\d:\d\d)Z<", core)
            if len(got) != 2 or any(t == "00:00:00" for _, _, t in got) or "Microsoft" not in app:
                bad.append(f)
    return bad


def pack_gates(out, meta, prompt_path, figs):
    files = sorted(os.listdir(out))
    fmts = sorted({os.path.splitext(f)[1] for f in files})
    C.ck("P01 input gate: 10 or more files", len(files) >= 10, len(files))
    C.ck("P02 input gate: 3 or more formats", len(fmts) >= 3, fmts)
    import pyarrow.parquet as pq
    n = pq.ParquetFile(os.path.join(out, Wr.SPINE)).metadata.num_rows
    C.ck("P03 input gate: a file of 25,000 or more rows", n >= 25000, n)
    C.ck("P04 input gate: two or more distractors, named in metadata.json, present, and off every solution path",
         len(meta["distractor_files"]) >= 2 and all(d in files for d in meta["distractor_files"])
         and not set(meta["distractor_files"]) & set(MAIN_FILES + DEVICE_FILES))
    leaks = [f for f in files if re.search(rb"(?i)distractor", open(os.path.join(out, f), "rb").read()) or "distractor" in f]
    C.ck("P05 the word distractor appears nowhere under target/", not leaks, leaks)
    texts = doc_texts(out)
    voc = [(f, m.group(0)) for f, t in texts.items() for m in re.finditer(r"(?i)\b(dispatch\w*|curtail\w*|4CP)\b", t)]
    voc += [(Wr.TERMS, "peak")] if re.search(r"(?i)peak", texts[Wr.TERMS]) else []
    near = [f for f, t in texts.items() if re.search(r"(?is)Business Saver.{0,200}\bpeak|\bpeak.{0,200}Business Saver", t)]
    C.ck("P06 programme vocabulary: no dispatch, curtail or 4CP anywhere, no 'peak' in the terms or near the "
         "programme's name", not voc and not near, (voc, near))
    hits = [(f, x) for f, t in texts.items() for x in figs if re.search(r"(?<![\d.,])" + re.escape(x) + r"(?![\d])", t)]
    C.ck("P07 no golden figure appears in any document", not hits, hits)
    em = [f for f, t in texts.items() if "\u2014" in t or "\u2013" in t]
    em += [f for f in files if f.endswith((".csv", ".md")) and "\u2014" in open(os.path.join(out, f)).read()]
    C.ck("P08 no em or en dash anywhere in the pack", not em, em)
    prompt = open(prompt_path).read()
    need = ["risk committee", "weather-zone books", "call options", "supply portfolio", "hedges"]
    banned = ["business saver", "harlan", "refrigerat", "cold", "replay", "percentile", "p90", "dispatch", "lender"]
    C.ck("P09 H20: the prompt carries the institutional nouns and none of the stump's words",
         all(x in prompt.lower() for x in need) and not any(x in prompt.lower() for x in banned))
    C.ck("P10 the deliverables the prompt names are not in the pack",
         not any(x in files for x in meta["deliverables"]))
    alltext = dict(texts)
    for f in files:
        if f.endswith(".csv"):
            alltext[f] = open(os.path.join(out, f)).read()
    homes = {k: sorted(f for f, t in alltext.items() if re.search(p, re.sub(r"\s+", " ", t))) for k, p in RULES.items()}
    C.ck("P11 single-statement invariant: each load-bearing rule stated in exactly one file",
         all(len(v) == 1 for v in homes.values()), homes)
    rows = {f: sum(1 for _ in open(os.path.join(out, f))) - 1 for f in files if f.endswith(".csv")}
    C.ck("P12 generation tell: no two CSVs share a row count", len(set(rows.values())) == len(rows), rows)
    bad = container_stamps(out)
    C.ck("P13 H1: every PDF and OOXML input carries a Central-time save stamp off midnight and an Office application name",
         not bad, bad)
    late = {}
    for f in files:
        if f.endswith(".csv"):
            ds = re.findall(r"\b(20\d\d-\d\d-\d\d)\b", open(os.path.join(out, f)).read())
            mx = max(ds) if ds else ""
            lim = "2027-12-31" if f == Wr.CALENDAR else "2028-12-31" if f == Wr.BLOTTER else "2027-04-09"
            if mx > lim:
                late[f] = mx
    C.ck("P14 H16: no extract carries a date after the 9 April 2027 extract (calendar and delivery periods excepted)",
         not late, late)
    logged = set(re.findall(r"`([^`]+)`", open(os.path.join(out, Wr.LOG)).read()))
    C.ck("P15 H9: the extract log names every shipped file but itself", logged == set(files) - {Wr.LOG},
         set(files) ^ logged)
    C.ck("P16 separation: no device file is on the main call's declared population", not set(MAIN_FILES) & set(DEVICE_FILES))
    stamps = {f: datetime.fromtimestamp(os.path.getmtime(os.path.join(out, f)), CT).isoformat() for f in files}
    C.ck("P17 every file's mtime is an in-fiction write time on or before 12 April 2027",
         all("2025-01-01" <= v[:10] <= "2027-04-12" for v in stamps.values()) and set(WRITTEN) == set(files), stamps)
    return {"files": files, "formats": fmts, "spine_rows": n, "rule_homes": homes, "csv_rows": rows}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    ap.add_argument("--meta", required=True)
    ap.add_argument("--record", default=None)
    ap.add_argument("--prompt", default=os.path.join(HERE, "..", "prompt.md"))
    a = ap.parse_args()

    w = pipeline.build_world()
    an = Analysis(w)
    main_res = C.main_call(an)
    grid_res = C.grid(an)
    corpus_res = C.corpus(an)
    twin_res = C.twins(an)
    C.clean_data(an)
    ask_res = C.asks(w)
    lend = C.lenders(an, D2.outlook_rows(), main_res["rungs"]["R0"])

    info = write_all(w, an, a.out)
    meta = {
        "task": "task124",
        "domain": "Business & Operations Analytics",
        "subdomain": "other (an energy retailer's summer supply planning)",
        "objective": "Forecasting & Predictive Modeling",
        "as_of": "2027-04-12",
        "deliverables": ["summer_block_committee.pptx", "summer_block_split.xlsx"],
        "distractor_files": list(Wr.DISTRACTORS),
        "distractor_notes": {
            Wr.DAYAHEAD: "Last summer's day-ahead forecasts of the book. The risk policy's exposure replays ten closed summers "
                         "at the coming book; no step reads a forecast of a closed summer.",
            Wr.OUTLOOK: "Wrong-basis distractor: ERCOT's zone outlook is the input to the lenders' zone-share basis, which "
                        "the risk policy (section 6) names as the lenders' view and not the exposure it defines. On that "
                        "basis the split is Coast-led and is neither the answer nor rung 0.",
        },
        "source": ("Every file is constructed for this task around a fictional Texas retail electricity provider "
                   "(Sabine Crest Energy), its customers, brokers and counterparties. ESI ID prefixes, ERCOT weather "
                   "zone and load profile codes, NAICS codes and NERC holidays follow the published conventions; load "
                   "series and the summer system peaks are synthetic, shaped on ERCOT's public hourly load by weather "
                   "zone, and are not a transcription. No real person, account or record is depicted."),
        "license": "CC BY 4.0",
        "created": "2026-10-10",
        "generator": "task124/generator/build_pack.py (seeded, deterministic)",
        "files": [],
    }
    for f in sorted(os.listdir(a.out)):
        p = os.path.join(a.out, f)
        meta["files"].append({"path": f, "format": os.path.splitext(f)[1][1:], "bytes": os.path.getsize(p)})
    figs = [f"{v:.1f}" for v in main_res["post"].values()] + [f"{main_res['loads']['North Central']:.1f}",
                                                                f"{main_res['centres_mw']:.1f}"]
    figs += [f"{int(round(v)):,}" for v in ask_res["premium"].values()] + [f"{v:.2f}" for v in ask_res["hedge_price"].values()]
    figs = sorted(set(figs))
    gates = pack_gates(a.out, meta, os.path.abspath(a.prompt), figs)
    with open(a.meta, "w") as f:
        json.dump(meta, f, indent=2)
        f.write("\n")
    C.ck("P18 metadata.json names the distractors and carries no golden figure",
         all(x not in json.dumps(meta) for x in figs) and meta["distractor_files"])
    record = {"main": main_res, "grid": grid_res, "corpus": {k: v for k, v in corpus_res.items()},
              "twins": twin_res, "asks": ask_res, "lenders": lend,
              "counts": {"amend_rows": an.amend_rows, "amend_kw": an.amend_kw, "centres_kw": an.centres,
                         "new_general_nc_kw": an.new_general_nc, "ded27": an.ded27, "rows27": an.rows27},
              "pack": gates, "info": info, "assertions": len(C.LOG)}
    for name, ok, _ in C.LOG:
        print(("PASS " if ok else "FAIL ") + name)
    print(f"\n{len(C.LOG)} assertions, all green. Split {main_res['rungs']['R5']}.")
    if a.record:
        with open(a.record, "w") as f:
            json.dump(record, f, indent=2, default=str)
            f.write("\n")


if __name__ == "__main__":
    main()
