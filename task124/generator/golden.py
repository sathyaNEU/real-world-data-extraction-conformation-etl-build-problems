"""Golden deliverables for task124 (Sabine Crest Energy, summer 2027 firm-capacity block).

Reads only the shipped bundle under target/, recomputes the split on the independent verifier's code path
(verify_pack.py, which never imports the generator), and writes the two files the prompt names into golden/:

    summer_block_committee.pptx   the split, the levelling chart, the cost of the block
    summer_block_split.xlsx       the exposure build, the replay, the lots, the premium and the hedge prices

    python3 task124/generator/golden.py [<target dir>] [<golden dir>]

Prints the critical components' figures. Every graded figure is asserted against the verifier's own checks.
"""
from __future__ import annotations

import contextlib
import io
import os
import sys
from datetime import datetime

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
sys.path.insert(0, HERE)

import verify_pack as V  # noqa: E402  (target-only code path)

BOOKS, NAMES = V.BOOKS, V.NAMES
NAME = dict(zip(BOOKS, NAMES))
MONTHS = [(6, "Jun-27"), (7, "Jul-27"), (8, "Aug-27"), (9, "Sep-27")]
STAMP = datetime(2027, 4, 14, 16, 42)
AUTHOR = "Tammy Ochoa"


def compute(target):
    """Run the verifier's main path quietly; refuse to write a golden if any of its checks fails."""
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        P = V.Pack(target)
        M = V.main_call(P)
        A = V.asks(target)
        V.corpus(P, M)  # the replay table: 80 of 80 under the per-kW replay
    bad = [n for n, ok, _ in V.RESULTS if not ok]
    assert not bad, bad
    yrs = P.years
    nc_md = M["book27"].groupby("book")["md"].sum().reindex(BOOKS).to_dict()
    base = dict(nc_md, NCENT=nc_md["NCENT"] - M["centres"] * 1000.0)
    extra = M["u"] * M["centres"]
    replay = {b: {y: M["pooled"][b][y] * base[b] / 1000.0 + (extra if b == "NCENT" else 0.0) for y in yrs}
              for b in BOOKS}
    load = {b: V.pct90(list(replay[b].values())) for b in BOOKS}
    hedges = M["hedges"]
    before = M["E"]["R4"]
    lots = M["L"]["R4"]
    after = M["post"]
    for b in BOOKS:
        assert abs(load[b] - hedges[b] - before[b]) < 1e-9
    assert sum(lots.values()) == 400
    assert V.vec(lots) == V.EXPECT["answer"]
    # the lot schedule, lot by lot (the policy's rule: each lot to the largest remaining uncovered exposure)
    sched, rem = [], dict(before)
    for k in range(80):
        b = max(BOOKS, key=lambda x: (rem[x], -BOOKS.index(x)))
        sched.append((k + 1, b, rem[b]))
        rem[b] -= 5
    assert all(abs(rem[b] - after[b]) < 1e-9 for b in BOOKS)
    nxt = max(BOOKS, key=lambda x: after[x])
    prem = {(b, m): round(A["premium"][f"{b}|{m}"]) for b in BOOKS for m, _ in MONTHS}
    hp = {b: round(A["hedge_price"][b], 2) for b in BOOKS}
    zones = zone_map(target)
    W = lambda x: int(round(x))  # noqa: E731  whole MW for every displayed megawatt figure
    # whole-MW display is internally consistent: load less hedges, and before less lots, in whole MW
    for b in BOOKS:
        assert W(load[b]) - hedges[b] == W(before[b]) and W(before[b]) - lots[b] == W(after[b]), b
    am = P.en[P.en["record_type"] == "AMEND"]
    amend = (len(am), am["md"].sum() / 1000.0)
    shift = lots["NCENT"] - M["L"]["R3"]["NCENT"]
    return dict(amend=amend, shift=shift, yrs=yrs, replay=replay, load=load, hedges=hedges, before=before, lots=lots, after=after,
                sched=sched, nxt=nxt, prem=prem, hp=hp, zones=zones, u=M["u"], cf=M["cf"], centres=M["centres"],
                n_centres=31, W=W)


def zone_map(target):
    import pandas as pd
    bm = pd.read_csv(os.path.join(target, "book_zone_map.csv"), dtype=str, keep_default_na=False)
    out = {}
    for b in BOOKS:
        r = bm[(bm["book"] == b) & (bm["effective_from"] <= "2027-06-01")
               & ((bm["effective_to"] == "") | (bm["effective_to"] >= "2027-09-30"))]
        assert len(r) == 1, b
        out[b] = r["load_zone"].iloc[0]
    return out


def split_title(G):
    order = sorted([b for b in BOOKS if G["lots"][b]], key=lambda b: -G["lots"][b])
    return "400 MW block: " + ", ".join(f"{NAME[b]} {G['lots'][b]}" for b in order) + ", none elsewhere"


# ----------------------------------------------------------------------------------------------- workbook
def write_xlsx(G, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    W = G["W"]
    wb = Workbook()
    bold, ital = Font(name="Calibri", bold=True, size=11), Font(name="Calibri", italic=True, size=9, color="595959")
    head_fill = PatternFill("solid", fgColor="DCE6F1")
    thin = Side(style="thin", color="8EA9C1")
    top = Border(top=thin)

    def header(ws, row, labels, widths=None):
        for j, lab in enumerate(labels, 1):
            c = ws.cell(row=row, column=j, value=lab)
            c.font, c.fill = bold, head_fill
            c.alignment = Alignment(wrap_text=True, vertical="bottom", horizontal="left" if j == 1 else "right")
        ws.row_dimensions[row].height = 30
        for j, w in enumerate(widths or [], 1):
            ws.column_dimensions[get_column_letter(j)].width = w

    # Split: one row per book, whole MW
    ws = wb.active
    ws.title = "Split"
    ws["A1"] = "Summer 2027 firm-capacity block: split by zone book"
    ws["A1"].font = Font(name="Calibri", bold=True, size=13)
    ws["A2"] = "Jun-Sep 2027 call options, 5 MW lots, 400 MW cap (RC minute 19 Mar 2027). Book as enrolled at the 9 Apr extract; hedges at close 9 Apr."
    ws["A2"].font = ital
    header(ws, 4, ["Book", "1-in-10 peak-hour load (MW)", "Summer hedges (MW)", "Uncovered before block (MW)",
                   "Lots (MW)", "Uncovered after block (MW)"], [16, 14, 12, 14, 10, 14])
    r = 5
    for b in BOOKS:
        vals = [NAME[b], W(G["load"][b]), G["hedges"][b], W(G["before"][b]), G["lots"][b], W(G["after"][b])]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            if j > 1:
                c.number_format = "#,##0"
        r += 1
    tot = ["Total", None, sum(G["hedges"].values()), None, sum(G["lots"].values()), None]
    for j, v in enumerate(tot, 1):
        c = ws.cell(row=r, column=j, value=v)
        c.font, c.border = bold, top
        if j > 1:
            c.number_format = "#,##0"
    ws.cell(row=r + 2, column=1, value=f"Next lot would go to {NAME[G['nxt']]} at {W(G['after'][G['nxt']])} MW uncovered.")
    ws.cell(row=r + 3, column=1, value="Loads and exposures carried unrounded through the levelling; shown in whole MW.").font = ital
    ws.freeze_panes = "B5"
    ws.print_area = f"A1:F{r + 3}"
    ws.page_setup.orientation = "landscape"

    # Replay: each closed summer's peak-hour coincidence at the 2027 book
    ws = wb.create_sheet("Replay")
    ws["A1"] = "Peak-hour load replayed at the 2027 enrolled book, MW (each summer's settled load at the ERCOT peak hour / that summer's enrolled max demand, times 2027 max demand)"
    ws["A1"].font = bold
    yrs = G["yrs"]
    header(ws, 3, ["Book"] + [str(y) for y in yrs] + ["P90 (1-in-10)"], [16] + [9] * len(yrs) + [12])
    r = 4
    for b in BOOKS:
        ws.cell(row=r, column=1, value=NAME[b])
        for j, y in enumerate(yrs, 2):
            ws.cell(row=r, column=j, value=round(G["replay"][b][y], 1)).number_format = "0.0"
        c = ws.cell(row=r, column=len(yrs) + 2, value=round(G["load"][b], 1))
        c.number_format, c.font = "0.0", bold
        r += 1
    notes = [
        "P90 by linear interpolation between closest ranks (risk policy s.2), ten closed summers 2017-2026.",
        f"North Central includes the {G['n_centres']} new Harlan Ridge distribution centres ({G['centres']:.1f} MW max demand) at "
        f"{G['u']:.2f} of max demand ({G['u'] * G['centres']:.1f} MW) in every summer.",
        f"0.90 is the cold-store members' draw at the peak hour on weekdays with no Business Saver window; at every closed "
        f"system peak they were inside a called window and drew {G['cf']:.2f}. New sites are not eligible before a metered summer.",
    ]
    for k, t in enumerate(notes):
        ws.cell(row=r + 1 + k, column=1, value=t).font = ital
    ws.freeze_panes = "B4"

    # Lots: the levelling, lot by lot
    ws = wb.create_sheet("Lots")
    header(ws, 1, ["Lot", "Book", "Uncovered before this lot (MW)", "Cumulative MW"], [6, 16, 16, 12])
    for k, b, e in G["sched"]:
        ws.append([k, NAME[b], round(e, 1), 5 * k])
        ws.cell(row=k + 1, column=3).number_format = "0.0"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:D{len(G['sched']) + 1}"
    ws.print_title_rows = "1:1"

    # Premium: lowest accepted quote per load zone and month, converted to USD per MW-month
    ws = wb.create_sheet("Premium")
    ws["A1"] = "Call option premium, USD per MW-month (lowest accepted quote from an approved counterparty, desk procedures s.3)"
    ws["A1"].font = bold
    header(ws, 3, ["Book", "Load zone"] + [lab for _, lab in MONTHS] + ["Lots (MW)", "Block premium (USD)"],
           [16, 13, 11, 11, 11, 11, 10, 15])
    r, total = 4, 0
    for b in BOOKS:
        row = [NAME[b], G["zones"][b]] + [G["prem"][(b, m)] for m, _ in MONTHS]
        cost = G["lots"][b] * sum(G["prem"][(b, m)] for m, _ in MONTHS)
        total += cost
        row += [G["lots"][b], cost]
        for j, v in enumerate(row, 1):
            c = ws.cell(row=r, column=j, value=v)
            if j > 2:
                c.number_format = "#,##0"
        r += 1
    c = ws.cell(row=r, column=1, value="Total")
    c.font = bold
    for j, v in ((7, sum(G["lots"].values())), (8, total)):
        c = ws.cell(row=r, column=j, value=v)
        c.font, c.border, c.number_format = bold, top, "#,##0"
    ws.cell(row=r + 2, column=1, value="Per-MWh quotes converted at the hours of the broker's notional in the month (s.2): 5x16 net of NERC holidays, or 7x16.").font = ital

    # Hedge price: average fixed price on existing summer strips
    ws = wb.create_sheet("Hedge price")
    ws["A1"] = "Existing Jun-Sep 2027 strips: average fixed price, USD/MWh (latest matched amendment, weighted by MWh in Jun-Sep, desk procedures s.4)"
    ws["A1"].font = bold
    header(ws, 3, ["Book", "Summer hedges (MW)", "Avg fixed price (USD/MWh)"], [16, 14, 16])
    for k, b in enumerate(BOOKS):
        ws.cell(row=4 + k, column=1, value=NAME[b])
        ws.cell(row=4 + k, column=2, value=G["hedges"][b]).number_format = "#,##0"
        ws.cell(row=4 + k, column=3, value=G["hp"][b]).number_format = "0.00"

    ws = wb.create_sheet("Notes")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110
    notes = [
        ("Prepared by", f"{AUTHOR}, Supply Portfolio, for the Risk Committee of 16 April 2027"),
        ("Book", "Enrolment extract of 9 April 2027, one row per ESI ID (an AMEND row supersedes the row it amends); every premise in service from 1 June"),
        ("Hedges", "Position report at close 9 April 2027, Summer 2027 tab, June to September strips (MW, same in every month)"),
        ("Peak hour", "ERCOT settled summer peak, hour ending CPT, 2017-2026; book loads from settlement at true-up"),
        ("1-in-10", "90th percentile of the ten replayed summers, inclusive interpolation (risk policy s.2)"),
        ("Lot rule", "5 MW lots, each to the largest remaining uncovered exposure, tie to the book listed first (risk policy s.5)"),
        ("Rounding", "MW carried unrounded through the levelling and shown in whole MW; premium to the dollar; hedge price to the cent"),
        ("Lenders' basis", "The zone-share view in the covenant pack (risk policy s.6) is the lenders' basis and is not used for the split"),
    ]
    for k, (a, b) in enumerate(notes, 1):
        ws.cell(row=k, column=1, value=a).font = bold
        ws.cell(row=k, column=2, value=b).alignment = Alignment(wrap_text=True, vertical="top")

    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.properties.created = wb.properties.modified = STAMP
    wb.properties.title = None
    wb.save(path)


# ----------------------------------------------------------------------------------------------- chart
BLUE, ORANGE, INK, INK2, GRID, SURF = "#2a78d6", "#eb6834", "#0b0b0b", "#52514e", "#e4e3df", "#fcfcfb"


def render_chart(G, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import rcParams
    rcParams.update({"font.family": "Carlito", "font.size": 11, "svg.hashsalt": "task124", "pdf.fonttype": 42})
    W = G["W"]
    order = sorted(BOOKS, key=lambda b: -G["before"][b])
    x = range(len(order))
    rest = [G["after"][b] for b in order]
    lots = [G["lots"][b] for b in order]
    fig, ax = plt.subplots(figsize=(11.4, 4.3), dpi=200)
    fig.patch.set_facecolor(SURF)
    ax.set_facecolor(SURF)
    ax.bar(x, rest, width=0.62, color=BLUE, edgecolor=SURF, linewidth=1.5, label="Uncovered after the block", zorder=2)
    ax.bar(x, lots, bottom=rest, width=0.62, color=ORANGE, edgecolor=SURF, linewidth=1.5, label="Bought in the block", zorder=2)
    for i, b in enumerate(order):
        top = G["before"][b]
        if G["lots"][b]:
            ax.text(i, top + 22, f"{W(top)}", ha="center", va="bottom", color=INK, fontsize=11, fontweight="bold")
            ax.text(i, top + 6, f"{G['lots'][b]} MW bought", ha="center", va="bottom", color=INK2, fontsize=9.5)
        else:
            ax.text(i, top - 6, f"{W(top)}", ha="center", va="top", color="white", fontsize=11, fontweight="bold")
    lvl, who = G["after"][G["nxt"]], NAME[G["nxt"]]
    ax.axhline(lvl, color=INK, linewidth=1.4, linestyle=(0, (5, 3)), zorder=3)
    ax.text(len(order) - 0.45, lvl + 32, f"Largest uncovered after the block: {W(lvl)} MW ({who})", ha="right",
            va="bottom", color=INK, fontsize=10.5)
    ax.set_xticks(list(x), [NAME[b] for b in order], color=INK)
    ax.set_ylabel("Uncovered 1-in-10 exposure before the block (MW)", color=INK2, fontsize=10)
    ax.set_ylim(0, 400)
    ax.yaxis.grid(True, color=GRID, linewidth=0.8, zorder=0)
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.spines["bottom"].set_color("#b5b4ae")
    ax.tick_params(axis="y", colors=INK2, length=0, labelsize=9.5)
    ax.tick_params(axis="x", length=0, labelsize=10.5)
    ax.set_title(split_title(G), loc="left", color=INK, fontsize=14, fontweight="bold", pad=12)
    ax.legend(loc="upper right", frameon=False, fontsize=9.5, ncol=2, bbox_to_anchor=(1.0, 1.0))
    fig.tight_layout()
    fig.savefig(path, facecolor=SURF, metadata={"Software": None})
    plt.close(fig)


# ----------------------------------------------------------------------------------------------- deck
def write_pptx(G, png, path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Inches, Pt

    W = G["W"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    ink, ink2, band = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E), RGBColor(0xDC, 0xE6, 0xF1)

    def text(slide, x, y, w, h, s, size=14, bold=False, color=ink):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        lines = s if isinstance(s, list) else [s]
        for k, ln in enumerate(lines):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.text = ln
            p.font.size, p.font.bold, p.font.color.rgb, p.font.name = Pt(size), bold, color, "Calibri"
            p.space_after = Pt(4)
        return tb

    def table(slide, x, y, w, rows, colw=None, size=11, h=0.3):
        t = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h * len(rows))).table
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                c = t.cell(i, j)
                c.text = str(v)
                p = c.text_frame.paragraphs[0]
                p.font.size, p.font.name = Pt(size), "Calibri"
                p.font.bold = i == 0 or (isinstance(r[0], str) and r[0] == "Total")
                p.font.color.rgb = ink
                c.fill.solid()
                c.fill.fore_color.rgb = band if i == 0 else RGBColor(0xFF, 0xFF, 0xFF)
                c.margin_top = c.margin_bottom = Emu(27432)
        if colw:
            for j, cw in enumerate(colw):
                t.columns[j].width = Inches(cw)
        return t

    def foot(slide, n):
        text(slide, 0.5, 7.05, 9.0, 0.3, "Sabine Crest Energy | Supply Portfolio | Risk Committee, 16 April 2027", 9, color=ink2)
        text(slide, 12.3, 7.05, 0.6, 0.3, str(n), 9, color=ink2)

    # 1. the split and the levelling chart
    s = prs.slides.add_slide(blank)
    nc_top = max(BOOKS, key=lambda b: G["lots"][b])
    text(s, 0.5, 0.3, 12.3, 0.6, f"{NAME[nc_top]} takes {G['lots'][nc_top]} MW of the 400 MW summer block; "
         f"{['no', 'one', 'two', 'three', 'four'][sum(1 for b in BOOKS if not G['lots'][b])]} books take none", 24, True)
    hdr = ["MW"] + NAMES + ["Total"]
    table(s, 0.5, 1.05, 12.3, [hdr, ["Lots"] + [G["lots"][b] for b in BOOKS] + [400]],
          colw=[1.1] + [1.24] * 8 + [1.28], size=13, h=0.36)
    s.shapes.add_picture(png, Inches(0.5), Inches(1.95), width=Inches(12.3))
    foot(s, 1)

    # 2. what the block costs
    s = prs.slides.add_slide(blank)
    total = sum(G["lots"][b] * sum(G["prem"][(b, m)] for m, _ in MONTHS) for b in BOOKS)
    text(s, 0.5, 0.3, 12.3, 0.6, f"The block costs ${total / 1e6:.2f} million in premium as split", 24, True)
    rows = [["Book", "Load zone", "Lots (MW)", "Jun-27", "Jul-27", "Aug-27", "Sep-27", "Existing summer hedges (MW)",
             "Avg hedge price ($/MWh)"]]
    for b in BOOKS:
        rows.append([NAME[b], G["zones"][b].replace("LZ_", "").title(), G["lots"][b]]
                    + [f"{G['prem'][(b, m)]:,}" for m, _ in MONTHS] + [G["hedges"][b], f"{G['hp'][b]:.2f}"])
    table(s, 0.5, 1.15, 12.3, rows, colw=[1.7, 1.2, 1.0, 1.15, 1.15, 1.15, 1.15, 2.1, 1.7], size=12, h=0.38)
    text(s, 0.5, 4.85, 12.3, 1.2, [
        "Premium in USD per MW-month at the desk's $250/MWh strike: lowest accepted quote from an approved counterparty "
        "for the book's 2027 load zone, per-MWh quotes converted at the hours of the broker's notional.",
        f"Premium for the 400 MW as split: ${total:,}. Hedge prices are the latest matched amendment, weighted by MWh over June to September.",
    ], 13, color=ink2)
    foot(s, 2)

    # 3. basis
    s = prs.slides.add_slide(blank)
    text(s, 0.5, 0.3, 12.3, 0.6, f"North Central carries {G['shift']} MW more because the new cold stores cannot be called", 24, True)
    text(s, 0.5, 1.2, 12.3, 5.5, [
        "Exposure per the 2027 risk policy: each book's load in the ERCOT summer peak hour at the 1-in-10 summer "
        "(P90 of the ten closed summers 2017-2026, inclusive), less its June to September strips. Lots go 5 MW at a time "
        "to the largest remaining exposure.",
        "Each closed summer is replayed at the book enrolled at the 9 April extract: that summer's settled peak-hour load over "
        "its enrolled maximum demand, applied to 2027 maximum demand. The same factors applied to the 2026 book reproduce all 80 "
        "cells of the replay table in last year's risk report.",
        "Enrolment amendment rows supersede the row they amend, so each premise counts once "
        f"({G['amend'][0]:,} North Central rows restate {G['amend'][1]:.0f} MW already on the book).",
        f"North Central takes on Harlan Ridge's {G['n_centres']} new cold-storage distribution centres, "
        f"{G['centres']:.1f} MW of maximum demand. Our fourteen cold stores and ice plants drew {G['cf']:.2f} of maximum demand at "
        f"every closed system peak, but every one of those peaks fell in a Business Saver window. On uncalled weekdays at the same "
        f"hour they draw {G['u']:.2f}.",
        f"The new centres cannot join Business Saver before a metered summer, so they are carried at {G['u']:.2f}: "
        f"{G['u'] * G['centres']:.1f} MW at the peak hour. Carrying them at the {G['cf']:.2f} the members drew instead would leave North Central {G['shift']} MW short.",
        f"After the block no book is left above {W(G['after'][G['nxt']])} MW uncovered; the next lot would go to {NAME[G['nxt']]}.",
    ], 15)
    foot(s, 3)

    prs.core_properties.author = AUTHOR
    prs.core_properties.last_modified_by = AUTHOR
    prs.core_properties.created = prs.core_properties.modified = STAMP
    prs.core_properties.title = "Summer 2027 block split"
    prs.core_properties.revision = 3
    prs.save(path)


# ----------------------------------------------------------------------------------------------- container
def set_props(path, application, title):
    """Rewrite docProps as the desktop app would have left them (author, in-fiction dates, no writer library) and pin
    every zip entry's time, so a rerun is byte-identical."""
    import re
    import zipfile
    from xml.sax.saxutils import escape
    iso = STAMP.strftime("%Y-%m-%dT%H:%M:%SZ")
    core = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
            '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
            'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
            'xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
            f'<dc:title>{escape(title)}</dc:title><dc:creator>{AUTHOR}</dc:creator>'
            f'<cp:lastModifiedBy>{AUTHOR}</cp:lastModifiedBy>'
            f'<dcterms:created xsi:type="dcterms:W3CDTF">{iso}</dcterms:created>'
            f'<dcterms:modified xsi:type="dcterms:W3CDTF">{iso}</dcterms:modified></cp:coreProperties>')
    app = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
           '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
           'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
           f'<Application>{application}</Application><DocSecurity>0</DocSecurity><ScaleCrop>false</ScaleCrop>'
           '<Company>Sabine Crest Energy</Company><LinksUpToDate>false</LinksUpToDate><SharedDoc>false</SharedDoc>'
           '<HyperlinksChanged>false</HyperlinksChanged><AppVersion>16.0000</AppVersion></Properties>')
    with zipfile.ZipFile(path) as z:
        names = [i.filename for i in z.infolist()]
        data = {n: z.read(n) for n in names}
    data["docProps/core.xml"], data["docProps/app.xml"] = core.encode(), app.encode()
    for n in names:
        if re.search(rb"(?i)python-pptx|openpyxl|matplotlib", data[n]) and not n.endswith((".png", ".jpeg")):
            raise AssertionError(f"writer signature left in {n}")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n in names:
            zi = zipfile.ZipInfo(n, date_time=STAMP.timetuple()[:6])
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, data[n])


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    target = args[0] if args else os.path.join(TASK, "target")
    out = args[1] if len(args) > 1 else os.path.join(TASK, "golden")
    scratch = os.environ.get("GOLDEN_SCRATCH", out)
    os.makedirs(out, exist_ok=True)
    G = compute(target)
    png = os.path.join(scratch, "_levelling_chart.png")
    render_chart(G, png)
    write_xlsx(G, os.path.join(out, "summer_block_split.xlsx"))
    write_pptx(G, png, os.path.join(out, "summer_block_committee.pptx"))
    set_props(os.path.join(out, "summer_block_split.xlsx"), "Microsoft Excel", "Summer 2027 block split")
    set_props(os.path.join(out, "summer_block_committee.pptx"), "Microsoft Office PowerPoint", "Summer 2027 block split")
    if scratch == out:
        os.remove(png)
    for f in ("summer_block_split.xlsx", "summer_block_committee.pptx"):
        os.utime(os.path.join(out, f), (STAMP.timestamp(), STAMP.timestamp()))
    W = G["W"]
    print("Split (MW):", " | ".join(f"{NAME[b]} {G['lots'][b]}" for b in BOOKS))
    print("North Central centres:", f"{G['n_centres']} sites, {G['centres']:.1f} MW max demand, uncalled draw {G['u']:.3f}, "
          f"called draw at closed peaks {G['cf']:.3f}, peak-hour draw {G['u'] * G['centres']:.1f} MW")
    print("1-in-10 load (MW):", " | ".join(f"{NAME[b]} {W(G['load'][b])}" for b in BOOKS))
    print("Uncovered before (MW):", " | ".join(f"{NAME[b]} {W(G['before'][b])}" for b in BOOKS))
    print("Uncovered after (MW):", " | ".join(f"{NAME[b]} {W(G['after'][b])}" for b in BOOKS))
    print("Next lot:", NAME[G["nxt"]], W(G["after"][G["nxt"]]), "MW")
    for b in BOOKS:
        print("Premium", NAME[b], [G["prem"][(b, m)] for m, _ in MONTHS], "hedge price", f"{G['hp'][b]:.2f}")


if __name__ == "__main__":
    main()
