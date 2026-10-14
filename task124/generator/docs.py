"""The pack's documents: the lead service line replacement plan (revision 2), the 2026 programme progress report,
the capital improvement plan board draft, the 2027 unit cost schedule, the 2026 change order log and the extract's
field notes."""
from __future__ import annotations

import datetime as dt

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import LETTER                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import inch                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import common as K   # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for _n, _f in (("LiberationSerif", "LiberationSerif-Regular.ttf"), ("LiberationSerif-Bold", "LiberationSerif-Bold.ttf"),
               ("LiberationSans", "LiberationSans-Regular.ttf"), ("LiberationSans-Bold", "LiberationSans-Bold.ttf")):
    pdfmetrics.registerFont(TTFont(_n, FONT_DIR + _f))
SERIF, SERIF_B, SANS, SANS_B = "LiberationSerif", "LiberationSerif-Bold", "LiberationSans", "LiberationSans-Bold"
INK = colors.HexColor("#1d2329")
GREY = colors.HexColor("#5d6670")
RULE = colors.HexColor("#9aa4ae")
BODY = ParagraphStyle("body", fontName=SERIF, fontSize=10.2, leading=13.6, textColor=INK, spaceAfter=5)
CLAUSE = ParagraphStyle("clause", parent=BODY, leftIndent=28, firstLineIndent=-28)
H1 = ParagraphStyle("h1", fontName=SANS_B, fontSize=14.5, leading=18, textColor=INK, spaceAfter=2)
H2 = ParagraphStyle("h2", fontName=SANS_B, fontSize=10.6, leading=14, textColor=INK, spaceBefore=8, spaceAfter=3)
META = ParagraphStyle("meta", fontName=SANS, fontSize=8.3, leading=10.8, textColor=GREY, spaceAfter=9)
CELL = ParagraphStyle("cell", fontName=SANS, fontSize=8.4, leading=10.2, textColor=INK)
CELLB = ParagraphStyle("cellb", parent=CELL, fontName=SANS_B)
CELLR = ParagraphStyle("cellr", parent=CELL, alignment=2)
CELLBR = ParagraphStyle("cellbr", parent=CELLB, alignment=2)
NOTE = ParagraphStyle("note", fontName=SANS, fontSize=7.8, leading=9.8, textColor=GREY, spaceBefore=3, spaceAfter=6)


def _frame(left, right):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont(SANS, 7.5)
        canvas.setFillColor(GREY)
        canvas.drawString(0.9 * inch, 0.55 * inch, left)
        canvas.drawRightString(LETTER[0] - 0.9 * inch, 0.55 * inch, "%s  |  Page %d" % (right, doc.page))
        canvas.restoreState()
    return draw


def _doc(path, title, author):
    return SimpleDocTemplate(path, pagesize=LETTER, leftMargin=0.9 * inch, rightMargin=0.9 * inch,
                             topMargin=0.8 * inch, bottomMargin=0.85 * inch, title=title, author=author,
                             subject="", creator="", producer="")


def _table(rows, widths, head=1, right=()):
    data = [[Paragraph(str(c), (CELLBR if j in right else CELLB) if i < head else (CELLR if j in right else CELL))
             for j, c in enumerate(r)] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, repeatRows=head)
    style = [("LINEBELOW", (0, head - 1), (-1, head - 1), 0.6, RULE), ("LINEBELOW", (0, -1), (-1, -1), 0.4, RULE),
             ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 2.2),
             ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2), ("LEFTPADDING", (0, 0), (-1, -1), 3)]
    t.setStyle(TableStyle(style))
    return t


def _money(x):
    return "{:,.0f}".format(x)


# --------------------------------------------------------------------------------------------- the plan, rev 2
def plan_clauses():
    return [
        ("1", "Purpose", [
            ("1.1", "This plan sets out how %s will replace every lead and galvanized-requiring-replacement (GRR) "
                    "service line in its distribution system, and the order in which it will work." % K.DISTRICT),
            ("1.2", "Under the Lead and Copper Rule Improvements the district must replace every lead and GRR service "
                    "line within ten years of the rule's compliance date of 1 November 2027, that is by 1 November "
                    "2037."),
        ]),
        ("2", "Scope", [
            ("2.1", "The plan covers every service line in the twelve replacement zones of section 3 that the service "
                    "line inventory classifies as Lead, GRR or Lead Status Unknown."),
            ("2.2", "Every such line is excavated at the curb stop and at the meter setting. A line found lead or GRR "
                    "is replaced in full, public and private side, at the same visit; a line found non-lead is restored "
                    "and reclassified in the inventory."),
            ("2.3", "A service line built in or after 1989 is non-lead by construction date under the district's 1988 "
                    "materials standard. Zones RZ-13 Darrow Meadows and RZ-14 Kittering North, built almost entirely "
                    "after 1988, are outside the plan: the lines on their few older streets are worked by district "
                    "crews under the operating budget in 2027 and 2028."),
        ]),
        ("3", "Replacement zones and priority order", [
            ("3.1", "The zones are worked in the priority order below, adopted by the board on 14 October 2026 on the "
                    "engineering ranking, which weighs each zone's schools and child care centres, its share of "
                    "premises built before 1930 and the state's disadvantaged community designation."),
        ]),
        ("4", "Plan cost", [
            ("4.1", "Plan costs are in 2027 dollars at the district's 2027 unit cost schedule, at the unit costs of the "
                    "zone in which the work is done."),
            ("4.2", "A zone's plan cost is its full replacements and its verification digs that find no lead, each at "
                    "the zone's unit cost, and one mobilisation, charged when the zone's work starts."),
            ("4.3", "The zone's Lead and GRR lines are full replacements. Each of its Lead Status Unknown lines is "
                    "carried as a full replacement at the share of lead or GRR expected among the district's Lead "
                    "Status Unknown lines of its construction era, and as a verification dig that finds no lead at "
                    "the rest."),
            ("4.4", "The construction eras are: built before 1930; built 1930 to 1949; built 1950 to 1988. A line's "
                    "construction year is the year built the inventory carries from the county assessor."),
        ]),
        ("5", "Schedule", [
            ("5.1", "The plan is funded from the lead service line construction line of the district's capital "
                    "improvement plan, one appropriation in each programme year from 2027."),
            ("5.2", "Zones are worked one after another in priority order, each zone's work starting when the zone "
                    "before it is finished. Work runs on from one zone into the next, and from one programme year into "
                    "the next, without waiting for a new programme year to start a zone."),
            ("5.3", "A zone is finished in the programme year in which the running total of appropriations first "
                    "covers the running total of plan cost through that zone."),
            ("5.4", "Each programme year's construction season runs from 1 April to 31 October. A zone finished in a "
                    "programme year is finished by 31 October of that year."),
        ]),
        ("6", "Inventory", [
            ("6.1", "The service line inventory is the district's record of each line's classification under the "
                    "Lead and Copper Rule Improvements (Lead, GRR, Non-lead, Lead Status Unknown) and of its basis: "
                    "the district's tap card, a field verification, the construction date, or a replacement. A line "
                    "built before 1989 with no usable tap card that has not been field verified is Lead Status "
                    "Unknown."),
            ("6.2", "Field verifications are recorded in the work management system (in the customer system's service "
                    "orders before 13 March 2023) and posted to the inventory."),
        ]),
    ]


def write_plan(path):
    doc = _doc(path, "Lead Service Line Replacement Plan, Revision 2", K.PEOPLE["engineer"])
    story = [Paragraph("Lead Service Line Replacement Plan", H1),
             Paragraph("%s. Revision 2, 20 October 2026. Owner: Engineering (%s, District Engineer). Supersedes "
                       "revision 1 of 30 September 2024." % (K.DISTRICT, K.PEOPLE["engineer"]), META)]
    for no, head, clauses in plan_clauses():
        story.append(Paragraph("%s&nbsp;&nbsp;%s" % (no, head), H2))
        for cno, text in clauses:
            story.append(Paragraph("%s&nbsp;&nbsp;&nbsp;%s" % (cno, text), CLAUSE))
        if no == "3":
            rows = [["Priority", "Zone", "Name", "Principal municipality"]]
            muni = {z[0]: (z[1], z[2]) for z in K.ZONES}
            for i, z in enumerate(K.PLAN_ORDER, 1):
                rows.append([str(i), z, muni[z][0], muni[z][1]])
            story.append(Spacer(1, 3))
            story.append(_table(rows, [0.75 * inch, 0.8 * inch, 1.8 * inch, 2.2 * inch]))
            story.append(Spacer(1, 4))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Approved for issue by %s, District Engineer, 20 October 2026." % K.PEOPLE["engineer"], META))
    doc.build(story, onFirstPage=_frame("%s  |  LSLR Plan, Revision 2" % K.DISTRICT, "20 October 2026"),
              onLaterPages=_frame("%s  |  LSLR Plan, Revision 2" % K.DISTRICT, "20 October 2026"))


# ------------------------------------------------------------------------------------ the 2026 progress report
def write_progress(path, t):
    """t: the report's tables, computed by the caller from the shipped files' own rows."""
    doc = _doc(path, "LSLR Programme 2026 Progress Report", K.PEOPLE["programme"])
    story = [Paragraph("Lead Service Line Replacement Programme: 2026 Progress Report", H1),
             Paragraph("Issued 6 January 2027 by the LSLR programme office (%s, Programme Manager). Covers January to "
                       "December 2026. This report counts work done; it does not schedule the plan or estimate its "
                       "cost. Version 1.0." % K.PEOPLE["programme"], META)]
    story.append(Paragraph("1&nbsp;&nbsp;The year", H2))
    all_2026 = t["replaced_by_year"].get(2026, 0)
    if all_2026 != t["replaced_year"]:
        done = ("replaced %s in full, %s of them under the three 2026 replacement contracts" % (
            "{:,}".format(all_2026), "{:,}".format(t["replaced_year"])))
    else:
        done = "replaced %s in full under the three 2026 replacement contracts" % "{:,}".format(t["replaced_year"])
    story.append(Paragraph(
        "District and contractor crews verified the material of %s service lines in 2026 and %s. The last "
        "verification of the season was on 18 December 2026; the 2027 season opens in April." % (
            "{:,}".format(t["verified_year"]), done), BODY))
    story.append(Paragraph("2&nbsp;&nbsp;Work by quarter", H2))
    rows = [["Quarter of 2026", "Service lines verified", "Full replacements completed"]]
    for q in (1, 2, 3, 4):
        rows.append(["Q%d" % q, "{:,}".format(t["verified_q"][q]), "{:,}".format(t["replaced_q"][q])])
    story.append(_table(rows, [1.6 * inch, 1.8 * inch, 2.1 * inch], right=(1, 2)))
    story.append(Paragraph("Quarters are calendar quarters in Eastern time. A line verified in two quarters is counted in "
                           "each. Most of the first quarter's replacements were lines found lead in the 2025 season.", NOTE))
    story.append(Paragraph("3&nbsp;&nbsp;Full replacements by year", H2))
    rows = [["Year"] + [str(y) for y in sorted(t["replaced_by_year"])],
            ["Full replacements"] + ["{:,}".format(t["replaced_by_year"][y]) for y in sorted(t["replaced_by_year"])]]
    story.append(_table(rows, [1.35 * inch] + [0.82 * inch] * len(t["replaced_by_year"]),
                        right=tuple(range(1, 1 + len(t["replaced_by_year"])))))
    story.append(Paragraph("4&nbsp;&nbsp;Inventory at 31 December 2026", H2))
    rows = [["Classification", "Service lines"]]
    for k in ("LEAD", "GRR", "NON-LEAD", "LEAD STATUS UNKNOWN"):
        rows.append([{"LEAD": "Lead", "GRR": "GRR", "NON-LEAD": "Non-lead",
                      "LEAD STATUS UNKNOWN": "Lead Status Unknown"}[k], "{:,}".format(t["inventory"][k])])
    rows.append(["All service lines", "{:,}".format(sum(t["inventory"].values()))])
    story.append(_table(rows, [2.4 * inch, 1.4 * inch], right=(1,)))
    story.append(Paragraph("System-wide, all fourteen zones. Lines replaced in full are Non-lead.", NOTE))
    story.append(Paragraph("5&nbsp;&nbsp;Crews", H2))
    story.append(Paragraph(
        "District service line crews SLV1 to SLV5 worked the full year. Brannigan Excavating (crews BRN-1 and BRN-2), "
        "Tri-County Underground (TCU-1 and TCU-2) and Kestrel Pipe and Utility (KPU-1) worked under the 2026 "
        "replacement contracts; Tri-County Underground's vendor record changed on 13 July 2026 after its "
        "acquisition, with its contract unchanged.", BODY))
    story.append(Paragraph("6&nbsp;&nbsp;Next", H2))
    story.append(Paragraph(
        "The 2027 to 2037 capital improvement plan goes to the board on 19 January 2027 with the replacement "
        "schedule for the twelve zones of plan revision 2.", BODY))
    doc.build(story, onFirstPage=_frame("%s  |  LSLR programme office" % K.DISTRICT, "2026 Progress Report"),
              onLaterPages=_frame("%s  |  LSLR programme office" % K.DISTRICT, "2026 Progress Report"))


# ------------------------------------------------------------------------------- the capital plan, board draft
CAPITAL_LINES = [
    ("Lead service line replacement, construction", "25.4", "2027 to 2037"),
    ("Lead service line programme, engineering, inspection and outreach", "2.3", "2027 to 2037"),
    ("Distribution system renewal", "11.8", "2027 to 2037"),
    ("Treatment plant filter and clearwell rehabilitation", "6.4", "2027 to 2030"),
    ("Storage tank recoating", "1.6", "2027 to 2033"),
    ("Meter and AMI replacement reserve", "0.9", "2031 to 2037"),
    ("Facilities and fleet", "1.2", "2027 to 2037"),
]


def write_capital(path):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    h = d.add_paragraph()
    run = h.add_run("Capital Improvement Plan 2027 to 2037")
    run.bold = True
    run.font.size = Pt(15)
    p = d.add_paragraph("%s. Board draft for adoption on 19 January 2027. Prepared by Finance (%s, Finance "
                        "Director), 14 December 2026." % (K.DISTRICT, K.PEOPLE["finance"]))
    p.runs[0].font.size = Pt(9)
    d.add_paragraph("The plan sets the district's capital programme for the eleven programme years 2027 to 2037. "
                    "Amounts are annual appropriations in 2027 dollars.")
    t = d.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for i, txt in enumerate(("Programme line", "Each year ($ million, 2027 dollars)", "Programme years")):
        t.rows[0].cells[i].text = txt
        t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for name, amt, yrs in CAPITAL_LINES:
        cells = t.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = name, amt, yrs
        cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.add_paragraph("")
    d.add_paragraph("The lead service line construction line funds contract construction, restoration and "
                    "mobilisation in the zones of the lead service line replacement plan; engineering, inspection, "
                    "outreach and district crews are carried on their own lines. Construction funds unspent in a "
                    "programme year carry forward to the next.")
    d.add_paragraph("Funding is from the 2027 and 2031 series of water revenue bonds, loans from the state drinking "
                    "water revolving fund and rates. The rate plan that accompanies this draft holds the residential "
                    "increase to 4.5 per cent a year through 2031.")
    d.add_paragraph("For adoption: the board is asked to adopt the plan and the 2027 appropriations on 19 January "
                    "2027. Chair: %s." % K.PEOPLE["chair"])
    d.core_properties.author = K.PEOPLE["finance"]
    d.core_properties.title = "Capital Improvement Plan 2027 to 2037"
    d.save(path)


# --------------------------------------------------------------------------------------- unit cost schedule
def _sheet_title(ws, lines):
    for i, (txt, bold, size) in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=txt).font = Font(bold=bold, size=size, italic=not bold and size < 10)


def _header(ws, row, cols):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font = Font(bold=True, size=10, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3A4D5C")
        cell.alignment = Alignment(wrap_text=True, vertical="top")


SYSTEM_RATES = (13_400, 2_450)
RATE_BASIS = {
    "RZ-02": "Downtown: pavement restoration and traffic control on both items",
    "RZ-04": "Long services on deep setbacks; curb stops under pavement",
    "RZ-07": "Short services, lawn restoration; dig restores sidewalk at the curb stop",
    "RZ-08": "System full replacement; dig restores sidewalk at the curb stop",
    "RZ-09": "System full replacement; dig restores sidewalk at the curb stop",
    "RZ-11": "System full replacement; curb stops under pavement",
}


def write_unit_costs(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2027 unit costs"
    _sheet_title(ws, [("%s, Engineering" % K.DISTRICT, True, 11),
                      ("Lead service line replacement: unit cost schedule, 2027 dollars", True, 11),
                      ("Engineer's estimate, 2 November 2026 (%s). Applies to work in programme year 2027." %
                       K.PEOPLE["engineer"], False, 9)])
    _header(ws, 5, ["Zone", "Zone name", "Full replacement, each ($)", "Verification dig, no lead found, each ($)",
                    "Rate basis"])
    r = 6
    for z in sorted(K.PLAN_ORDER):
        uf, ud = K.ZONE_PRICES[z]
        ws.cell(row=r, column=1, value=z)
        ws.cell(row=r, column=2, value=K.ZONE_NAME[z])
        ws.cell(row=r, column=3, value=uf).number_format = "#,##0"
        ws.cell(row=r, column=4, value=ud).number_format = "#,##0"
        ws.cell(row=r, column=5, value=RATE_BASIS.get(z, "System rates" if (uf, ud) == SYSTEM_RATES else None))
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Mobilisation, each zone ($)").font = Font(bold=True)
    ws.cell(row=r, column=3, value=K.MOBILISATION).number_format = "#,##0"
    r += 2
    for txt in ("Full replacement: public and private side from the main to the meter setting, new curb stop and "
                "setter, restoration.",
                "Verification dig: excavation at the curb stop and at the meter setting, restoration; no replacement.",
                "System rates: $%s a full replacement and $%s a verification dig. A zone's rate differs from the "
                "system rate where its typical service length, pavement or traffic control differs (column E)." % (
                    "{:,}".format(SYSTEM_RATES[0]), "{:,}".format(SYSTEM_RATES[1]))):
        ws.cell(row=r, column=1, value=txt).font = Font(size=9)
        r += 1
    for j, w in enumerate((10, 20, 16, 20, 46), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[5].height = 30
    wb.save(path)


# --------------------------------------------------------------------------------------- change order log
def write_change_orders(path, co, contractor_of):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CO log 2026"
    _sheet_title(ws, [("%s, Engineering: change order log, 2026 lead service line replacement contracts" % K.DISTRICT,
                       True, 11),
                      ("Approved change orders, paid to the contractor on its pay applications. Each change order is "
                       "entered once, with the sites it covers.", False, 9)])
    cols = ["CO no.", "Contract", "Contractor", "Approved", "Description", "Sites covered", "No. of sites",
            "Amount ($)"]
    _header(ws, 4, cols)
    r = 5
    for c in co.sort_values("co_no").itertuples():
        vals = [c.co_no, c.contract, contractor_of[c.contract], c.approved, c.description, ", ".join(c.sites),
                len(c.sites), c.amount]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=j, value=v)
            if j == 4:
                cell.number_format = "yyyy-mm-dd"
            if j == 8:
                cell.number_format = "#,##0.00"
        r += 1
    for j, w in enumerate((11, 12, 24, 11, 44, 40, 8, 12), 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A5"
    wb.save(path)


# ------------------------------------------------------------------------------------------------- field notes
def write_notes(path, counts):
    inv = K.INVENTORY
    text = """# LSL data extract, January 2027: field notes

Prepared by GIS and Asset Records for the LSLR programme office, 8 January 2027. One section per file. The extract
log ({log}) gives each file's source system, date and row count.

## {inv}

One row per service line (main to meter setting) in the district's lead service line inventory, all fourteen zones,
as of 8 January 2027.

- `sl_id`: GIS service line id.
- `service_address`, `municipality`: the premise address.
- `zone`: the line's replacement zone under plan revision 2, RZ-01 to RZ-14.
- `main_segment`: GIS id of the water main segment the service taps.
- `year_built`: the premise's year built, from the county assessor; 0 where the assessor carries none (most
  non-residential parcels).
- `tap_card_material`: the material on the district's tap card (LEAD, GALV, COPPER, PLASTIC); blank where no usable
  card survives.
- `sl_class`: LEAD, GRR, NON-LEAD or LEAD STATUS UNKNOWN.
- `class_basis`: RECORDS (tap card), FIELD VERIFIED, CONSTRUCTION DATE (built in or after 1989) or REPLACED (found
  lead or GRR and replaced in full, NON-LEAD since); blank for Lead Status Unknown. A galvanized tap card line is GRR
  until it is verified.
- `field_verified_on`: date of the line's latest field verification.
- `replaced_on`: date the line was replaced in full.

## {ver}

One row per field verification of a service line's material, 4 January 2021 to 18 December 2026. A line verified
more than once has a row for each visit.

- `source`: CIS for the customer system's service orders (to 12 March 2023), WMS for the work management system
  (from 13 March 2023).
- `wo_no`: service order (CIS) or work order (WMS) number.
- `verified_date`: MM/DD/YYYY on CIS rows, YYYY-MM-DD on WMS rows.
- `premise_no`: customer system premise number, CIS rows.
- `sl_id`: GIS service line id, WMS rows.
- `method`: EXCAVATION, POTHOLE (vacuum excavation at the curb stop), CURB BOX CAMERA, METER PIT VISUAL.
- `crew`: district crews SLV1 to SLV5; contractor crews BRN-1 and BRN-2 (Brannigan Excavating), TCU-1 and TCU-2
  (Tri-County Underground), KPU-1 (Kestrel Pipe and Utility).
- `result_code`: the material found, in the codes of the form it was recorded on ({codes}).

## {xw}

Customer system premise numbers to GIS service line ids, built for the 13 March 2023 cut-over. `legacy_premise_no`
is zero-padded to eight digits; the customer system printed premise numbers without the leading zeros.
`match_method` is how the pair was matched (ADDRESS, METER, FIELD).

## {codes}

Every result code used on the three recording routes, with the material and the class it posts to the inventory.

## {comp}

Full replacement completions under the 2026 replacement contracts, from the contract management module of the work
management system. One row per closing of a work order.

- `site_id`: GIS service line id.
- `completed_utc`: completion time stamped by the field tablet, in UTC.
- `closed_on`: date the work order was closed.
- `row_status`: CLOSED, or RECLOSED when a closed work order was reopened for a correction and closed again; the
  export keeps both closings, and the reclosing's paid amounts replace the first closing's.
- `paid_base`, `paid_extra_ft`, `paid_restoration`: paid on the pay application for the site: the bid price, the
  service length past 50 feet at the contract rate, and restoration.
- `co_no`, `co_amount`: the approved change order that covers the site, and the change order's amount.

## {con}

The three 2026 replacement contracts, with the vendor records they were paid under, award dates, zones and bid
prices.

## {co}

The engineer's log of approved change orders on the 2026 contracts.

## {ag}

Right-of-entry agreements for work on the private side of the line, from the agreement system, one row per form.
Agreements on lines since replaced or verified non-lead are archived and are not in this file.

- `signed_on`: date the form was signed.
- `channel`: PAPER (mailed form) or PORTAL (online, from 3 June 2024).
- `premise_no`: customer system premise number, on forms signed before 13 March 2023.
- `sl_id`: GIS service line id, on forms signed from 13 March 2023.
- `zone_on_form`: the replacement zone printed on the form.
- `status`, `status_date`: SIGNED or WITHDRAWN, and the date of that status.

## {out}

The programme office's outreach tracker: one row per line open in the plan (Lead, GRR or Lead Status Unknown in the
twelve zones of plan revision 2).

- `zone`: the line's replacement zone under plan revision 2.
- `letter_1_mailed`: the date the first outreach letter for the line went to its owner.
- `portal_invite_sent`: the date the owner was invited to the right-of-entry form on the customer portal.
- `portal_status`, `portal_status_date`: the portal's status for the line and its date: NOT OPENED, OPENED, STARTED,
  E-SIGNED (the owner signed the right-of-entry form on the portal) or DECLINED.

## Documents

- {plan}: the replacement plan, revision 2.
- {cap}: the capital improvement plan, board draft.
- {uc}: the 2027 unit cost schedule.
- {prog}: the programme office's 2026 progress report.

## Other extracts in the folder

- {tap}: lead and copper compliance samples at the tier 1 sites, 2025 and 2026, from the laboratory system.
- {mx}: meter exchanges under the automated metering rollout, April 2023 to December 2025, from the meter shop.
""".format(log=K.EXTRACT_LOG, inv=inv, ver=K.VERIFICATIONS, xw=K.CROSSWALK, codes=K.CODES, comp=K.COMPLETIONS,
           con=K.CONTRACTS, co=K.CHANGE_ORDERS, ag=K.AGREEMENTS, out=K.OUTREACH, plan=K.PLAN, cap=K.CAPITAL,
           uc=K.UNIT_COSTS,
           prog=K.PROGRESS, tap=K.TAP_SAMPLES, mx=K.METER_EXCHANGES)
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(text)
