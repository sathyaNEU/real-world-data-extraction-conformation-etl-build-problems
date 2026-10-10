"""Word and Excel documents of the pack, and the two markdown files."""
from __future__ import annotations

import re
import zipfile
from datetime import date, datetime, timezone
from xml.sax.saxutils import escape

import desk as K
from common import BOOKS, CODE, SUMMERS


# ------------------------------------------------------------------------------ Word
def docx_stats(texts, cells, line_chars=92, cell_chars=22, page_lines=40):
    flat = [t for t in texts if t.strip()] + [c for row in cells for c in row if c.strip()]
    words = sum(len(t.split()) for t in flat)
    chars = sum(len(re.sub(r"\s", "", t)) for t in flat)
    spaced = sum(len(t) for t in flat)
    lines = sum(-(-len(t) // line_chars) for t in texts if t.strip())
    lines += sum(max(-(-len(c) // cell_chars) for c in row) for row in cells) if cells else 0
    return {"words": words, "chars": chars, "spaced": spaced, "paragraphs": len(flat), "lines": lines,
            "pages": 1 + (lines - 1) // page_lines}


def finish_docx(path, texts, cells, title, saved):
    """Package properties as Word writes them on save, and no preview picture."""
    st = docx_stats(texts, cells)
    iso = lambda d: d.strftime("%Y-%m-%dT%H:%M:%SZ")  # noqa: E731
    org = escape("Sabine Crest Energy LLC")
    core = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
            '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
            'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
            'xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
            f'<dc:title>{escape(title)}</dc:title><dc:subject></dc:subject><dc:creator>{escape(saved["author"])}</dc:creator>'
            f'<cp:keywords></cp:keywords><dc:description></dc:description>'
            f'<cp:lastModifiedBy>{escape(saved["by"])}</cp:lastModifiedBy><cp:revision>{saved["revision"]}</cp:revision>'
            f'<dcterms:created xsi:type="dcterms:W3CDTF">{iso(saved["created"])}</dcterms:created>'
            f'<dcterms:modified xsi:type="dcterms:W3CDTF">{iso(saved["modified"])}</dcterms:modified>'
            '</cp:coreProperties>')
    app = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
           '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
           'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
           f'<Template>Normal.dotm</Template><TotalTime>{saved["minutes"]}</TotalTime>'
           f'<Pages>{st["pages"]}</Pages><Words>{st["words"]}</Words><Characters>{st["chars"]}</Characters>'
           '<Application>Microsoft Office Word</Application><DocSecurity>0</DocSecurity>'
           f'<Lines>{st["lines"]}</Lines><Paragraphs>{st["paragraphs"]}</Paragraphs><ScaleCrop>false</ScaleCrop>'
           '<HeadingPairs><vt:vector size="2" baseType="variant"><vt:variant><vt:lpstr>Title</vt:lpstr></vt:variant>'
           '<vt:variant><vt:i4>1</vt:i4></vt:variant></vt:vector></HeadingPairs><TitlesOfParts>'
           f'<vt:vector size="1" baseType="lpstr"><vt:lpstr>{escape(title)}</vt:lpstr></vt:vector></TitlesOfParts>'
           f'<Company>{org}</Company><LinksUpToDate>false</LinksUpToDate>'
           f'<CharactersWithSpaces>{st["spaced"]}</CharactersWithSpaces><SharedDoc>false</SharedDoc>'
           '<HyperlinksChanged>false</HyperlinksChanged><AppVersion>16.0000</AppVersion></Properties>')
    with zipfile.ZipFile(path) as z:
        infos = [i for i in z.infolist() if i.filename != "docProps/thumbnail.jpeg"]
        data = {i.filename: z.read(i.filename) for i in infos}
    data["docProps/core.xml"] = core.encode()
    data["docProps/app.xml"] = app.encode()
    rels = data["_rels/.rels"].decode()
    rels, n = re.subn(r'<Relationship Id="rId\d+" Type="[^"]*/metadata/thumbnail" Target="docProps/thumbnail\.jpeg"/>',
                      "", rels)
    assert n == 1, rels
    data["_rels/.rels"] = rels.encode()
    ct = data["[Content_Types].xml"].decode()
    ct, n = re.subn(r'<Default Extension="jpeg" ContentType="image/jpeg"/>', "", ct)
    assert n == 1, ct
    data["[Content_Types].xml"] = ct.encode()
    stg = data["word/settings.xml"].decode()
    stg, n = re.subn(r"<w:savePreviewPicture/>", "", stg)
    assert n == 1
    data["word/settings.xml"] = stg.encode()
    dt = saved["modified"].astimezone(timezone.utc)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for i in infos:
            zi = zipfile.ZipInfo(i.filename, date_time=(dt.year, dt.month, dt.day, dt.hour, dt.minute, dt.second))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            z.writestr(zi, data[i.filename])
    return st


def _new_doc():
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)

    def p(text, bold=False, size=None, space=4, italic=False):
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        para.paragraph_format.space_after = Pt(space)
        return para
    return doc, p


def _texts(doc):
    return [x.text for x in doc.paragraphs], [[c.text for c in row.cells] for t in doc.tables for row in t.rows]


POLICY_SAVED = {"author": "Tammy Ochoa", "by": "Tammy Ochoa", "revision": 9, "minutes": 212,
                "created": datetime(2027, 1, 26, 15, 4, tzinfo=timezone.utc),
                "modified": datetime(2027, 2, 19, 20, 41, tzinfo=timezone.utc)}
PROC_SAVED = {"author": "Gregory Sheppard", "by": "Gregory Sheppard", "revision": 14, "minutes": 365,
              "created": datetime(2024, 11, 4, 14, 12, tzinfo=timezone.utc),
              "modified": datetime(2027, 1, 15, 22, 7, tzinfo=timezone.utc)}


def policy(path):
    doc, p = _new_doc()
    p("SABINE CREST ENERGY LLC", bold=True, size=9, space=0)
    p("Summer Supply Risk Policy", bold=True, size=14, space=2)
    p("Adopted by the Risk Committee on 19 February 2027 for the 2027 summer. Owner: Supply Portfolio "
      "(Tammy Ochoa). Approved: Susan Kelley, Chair of the Risk Committee.", size=9, space=10)
    p("1. Scope", bold=True)
    p("This policy governs the summer firm capacity Sabine Crest buys for its retail book. The book is run as eight "
      "weather-zone books, listed in this order: Coast, East, Far West, North, North Central, South Central, "
      "Southern and West.")
    p("2. Uncovered exposure", bold=True)
    p("A book's uncovered exposure is its load in the system peak hour of the coming summer at the 1-in-10 summer, "
      "less its summer hedges. The 1-in-10 summer is the 90th percentile across the ten most recent closed "
      "summers, by linear interpolation between closest ranks (the inclusive method). Summer hedges are the book's "
      "June to September strips in the position report taken with the spring enrolment extract, in MW.")
    p("3. The coming book", bold=True)
    p("The coming book is the book enrolled at the spring enrolment extract. Every premise on it is taken to be in "
      "service from 1 June at its enrolled maximum demand.")
    p("4. System peak hour", bold=True)
    p("The system peak hour of a summer is ERCOT's settled summer peak, hour ending, Central Prevailing Time, as "
      "ERCOT publishes it. Book loads are taken from ERCOT settlement at true-up.")
    p("5. The summer block", bold=True)
    p("The Risk Committee caps the summer block each spring. Supply Portfolio places the block in 5 MW lots: each "
      "lot goes to the book with the largest remaining uncovered exposure, a tie going to the book listed first in "
      "section 1. The split is adopted by the Committee before the options are bought.")
    p("6. Lenders' view", bold=True)
    p("Our lenders review summer exposure on the zone-share basis: each book's share of its weather zone's June to "
      "September energy in the last closed summer, applied to ERCOT's published 1-in-10 summer peak for the zone. "
      "Treasury will put that view in the covenant pack in May and it is what the lenders will have in front of "
      "them. It is their basis, not the exposure this policy defines.")
    p("7. Review", bold=True)
    p("Supply Portfolio reviews this policy each winter with Load Planning and Trading.", space=10)
    p("Internal. Do not forward outside Sabine Crest.", size=8, italic=True)
    title = "Summer Supply Risk Policy 2027"
    doc.core_properties.title = title
    doc.save(path)
    t, c = _texts(doc)
    finish_docx(path, t, c, title, POLICY_SAVED)


def procedures(path):
    from docx.shared import Pt
    doc, p = _new_doc()
    p("SABINE CREST ENERGY | TRADING", bold=True, size=9, space=0)
    p("Desk Procedures: ERCOT summer options and hedge reporting", bold=True, size=13, space=2)
    p("Rev. 3, January 2027. Approved: Gregory Sheppard, Head of Trading.", size=9, space=10)
    p("1. Approved brokers", bold=True)
    p("Summer call options are bought through the three approved brokers below. Each quotes its premium in its own "
      "convention.")
    tab = doc.add_table(rows=1, cols=2)
    tab.style = "Table Grid"
    for c, t in zip(tab.rows[0].cells, ["Broker", "Premium quoted in"]):
        c.text = t
    for b in ("GEB", "TBC", "PPB"):
        cells = tab.add_row().cells
        cells[0].text = K.BROKERS[b][0]
        cells[1].text = f"USD per MWh of {K.BROKERS[b][1]} notional"
    for row in tab.rows:
        for c in row.cells:
            for para in c.paragraphs:
                for r in para.runs:
                    r.font.size = Pt(9.5)
    p("")
    p("2. Converting a premium", bold=True)
    p("A premium quoted per MWh is converted to USD per MW-month at the hours of the notional it is quoted on in the "
      "delivery month. 5x16 is hours ending 7 through 22, Monday to Friday, other than NERC holidays; 7x16 is hours "
      "ending 7 through 22 every day.")
    p("3. Choosing a quote", bold=True)
    p("Calls are bought at the desk's standard strike of $250/MWh, daily exercise. The desk records a decision on "
      "every quote in the quote decisions log. Each lot is bought at the lowest accepted quote for its book's load "
      "zone and delivery month; a book's load zone for a month is the one the book map gives for that month. A quote "
      "counts only if the option writer was an approved counterparty on the day the quote was sent.")
    p("4. Hedge reporting", bold=True)
    p("The version of record of a trade is its latest amendment matched to the counterparty's confirmation. A "
      "book's average fixed price for a period weights each trade by its MWh in the period. Portfolios roll up to "
      "books through the portfolio crosswalk.")
    p("5. Records", bold=True)
    p("Trade Operations keeps the blotter, the confirmation matching log and the counterparty master. Credit keeps "
      "approvals current.", space=10)
    p("Internal. Trading desk use.", size=8, italic=True)
    title = "Desk Procedures, ERCOT summer options and hedge reporting"
    doc.core_properties.title = title
    doc.save(path)
    t, c = _texts(doc)
    finish_docx(path, t, c, title, PROC_SAVED)


# ------------------------------------------------------------------------------ Excel
def _wb(path, title, author, company, created):
    import xlsxwriter
    wb = xlsxwriter.Workbook(path)
    wb.set_properties({"title": title, "author": author, "company": company, "created": created})
    return wb


def positions(path, w):
    bl = w.blotter
    s = bl[(bl["start"] == date(2027, 6, 1)) & (bl["amend"] == 0)].copy()
    pf = {p: b for b, ps in K.PORTFOLIOS.items() for p in ps}
    s["book"] = s["portfolio"].map(pf)
    s["shape"] = s["product"].str.split(".").str[2].str.lower()
    agg = s.groupby(["book", "portfolio", "shape"])["mw"].sum()
    wb = _wb(path, "Position report", "Trade Operations", "Sabine Crest Energy LLC", datetime(2027, 4, 9, 23, 18, 0))
    ws = wb.add_worksheet("Summer 2027")
    bold = wb.add_format({"bold": True})
    hdr = wb.add_format({"bold": True, "bottom": 1})
    num = wb.add_format({"num_format": "#,##0"})
    tot = wb.add_format({"bold": True, "num_format": "#,##0", "top": 1})
    ws.write(0, 0, "Sabine Crest Energy, ERCOT retail hedge book: position report", bold)
    ws.write(1, 0, "As of close 9 April 2027. Fixed-price strips by delivery month, MW. Prices are in the blotter.")
    heads = ["Book", "Portfolio", "Shape", "Jun-27", "Jul-27", "Aug-27", "Sep-27"]
    for j, h in enumerate(heads):
        ws.write(3, j, h, hdr)
    i = 4
    for b in BOOKS:
        for (bb, pfo, shp), mw in agg.items():
            if bb != b:
                continue
            ws.write(i, 0, CODE[b]); ws.write(i, 1, pfo); ws.write(i, 2, shp.upper())
            for j in range(4):
                ws.write_number(i, 3 + j, int(mw), num)
            i += 1
        ws.write(i, 0, f"{CODE[b]} total", bold)
        for j in range(4):
            ws.write_number(i, 3 + j, int(agg[b].sum()), tot)
        i += 2
    ws.set_column(0, 0, 13); ws.set_column(1, 1, 10); ws.set_column(2, 6, 8)
    ws.freeze_panes(4, 0)
    ws2 = wb.add_worksheet("Q4-27 and Cal-28")
    ws2.write(0, 0, "Fixed-price strips delivering after September 2027, MW", bold)
    o = bl[(bl["start"] > date(2027, 9, 30)) & (bl["amend"] == 0)].copy()
    o["book"] = o["portfolio"].map(pf)
    for j, h in enumerate(["Book", "Portfolio", "Product", "Delivery start", "Delivery end", "MW"]):
        ws2.write(2, j, h, hdr)
    for k, r in enumerate(o.sort_values(["start", "trade_id"]).itertuples(), start=3):
        ws2.write(k, 0, CODE[r.book]); ws2.write(k, 1, r.portfolio); ws2.write(k, 2, r.product)
        ws2.write(k, 3, r.start.isoformat()); ws2.write(k, 4, r.end.isoformat()); ws2.write_number(k, 5, int(r.mw), num)
    ws2.set_column(0, 2, 14); ws2.set_column(3, 4, 12)
    wb.close()
    return agg


def roster(path, w):
    p = w.prem[(w.prem["comp"] == "ref") & (w.prem["record_type"] == "NEW")]
    g = p.groupby(["account", "customer"]).agg(sites=("esi_id", "count"), since=("start", "min")).reset_index()
    g = g.sort_values("account")
    wb = _wb(path, "Business Saver roster 2027", "Donald Lee", "Sabine Crest Energy LLC", datetime(2027, 3, 31, 21, 46, 0))
    ws = wb.add_worksheet("2027 season")
    bold = wb.add_format({"bold": True})
    hdr = wb.add_format({"bold": True, "bottom": 1})
    ws.write(0, 0, "Business Saver: member accounts for the 2027 season", bold)
    ws.write(1, 0, "Program Desk. Renewals closed 31 March 2027.")
    for j, h in enumerate(["Account", "Customer", "Sites", "Member since", "2027 status"]):
        ws.write(3, j, h, hdr)
    for i, r in enumerate(g.itertuples(), start=4):
        ws.write(i, 0, r.account); ws.write(i, 1, r.customer); ws.write_number(i, 2, int(r.sites))
        ws.write_number(i, 3, max(2015, int(r.since.year))); ws.write(i, 4, "Renewed")
    ws.set_column(0, 0, 12); ws.set_column(1, 1, 40); ws.set_column(2, 4, 12)
    wb.close()


def pecos(path, w):
    q = w.quotes[w.quotes["broker"] == "PPB"].sort_values(["sent", "qid"], kind="mergesort")
    wb = _wb(path, "Sabine Crest - ERCOT summer calls", "Pecos Power Brokerage", "Pecos Power Brokerage LLC",
             datetime(2027, 4, 9, 20, 2, 0))
    ws = wb.add_worksheet("Offers")
    bold = wb.add_format({"bold": True})
    hdr = wb.add_format({"bold": True, "bg_color": "#D9E1F2", "border": 1})
    money = wb.add_format({"num_format": "0.00"})
    ws.write(0, 0, "PECOS POWER BROKERAGE", bold)
    ws.write(1, 0, "Firm offers to Sabine Crest Energy: ERCOT monthly calls, strike $250/MWh, daily exercise, 5x16")
    for j, h in enumerate(["Ref", "Seller", "Zone", "Month", "Strike", "Premium $/MWh", "Sent (CPT)"]):
        ws.write(3, j, h, hdr)
    for i, r in enumerate(q.itertuples(), start=4):
        ws.write(i, 0, r.qid); ws.write(i, 1, r.seller); ws.write(i, 2, K.PECOS_ZONE[r.zone])
        ws.write(i, 3, date(2027, r.month, 1).strftime("%b-%y")); ws.write_number(i, 4, 250, money)
        ws.write_number(i, 5, float(r.price), money); ws.write(i, 6, r.sent.strftime("%m/%d/%Y %I:%M %p"))
    ws.set_column(0, 0, 13); ws.set_column(1, 1, 30); ws.set_column(2, 3, 11); ws.set_column(4, 5, 13)
    ws.set_column(6, 6, 20)
    wb.close()


ZONE_OUTLOOK = {  # 2026 summer peak MW, 2026 June-September load factor, 2027 50/50 growth
    "Coast": (23_412, 0.712, 0.031), "East": (3_086, 0.668, 0.022), "Far West": (8_964, 0.861, 0.074),
    "North": (1_752, 0.664, 0.012), "North Central": (28_655, 0.693, 0.038), "South Central": (15_318, 0.671, 0.041),
    "Southern": (7_094, 0.703, 0.027), "West": (2_247, 0.719, 0.019)}


def outlook_rows():
    hrs = 122 * 24
    out = []
    for b in BOOKS:
        pk, lf, g = ZONE_OUTLOOK[b]
        e = round(pk * lf * hrs / 1000.0)
        f50 = round(pk * (1 + g))
        f90 = round(f50 * 1.047)
        out.append((b, pk, e, f50, f90))
    return out


def outlook(path):
    wb = _wb(path, "Summer 2027 weather zone outlook", "ERCOT", "Electric Reliability Council of Texas",
             datetime(2026, 12, 18, 16, 30, 0))
    ws = wb.add_worksheet("Weather zones")
    bold = wb.add_format({"bold": True})
    hdr = wb.add_format({"bold": True, "bottom": 1, "text_wrap": True})
    num = wb.add_format({"num_format": "#,##0"})
    ws.write(0, 0, "Summer 2027 load outlook by weather zone", bold)
    ws.write(1, 0, "Weather zone peaks are coincident with the zone's own summer peak hour. Energy is June through "
                   "September.")
    heads = ["Weather zone", "2026 summer peak (MW)", "2026 Jun-Sep energy (GWh)", "2027 peak forecast, 50/50 (MW)",
             "2027 peak forecast, 90/10 (MW)"]
    for j, h in enumerate(heads):
        ws.write(3, j, h, hdr)
    for i, (b, pk, e, f50, f90) in enumerate(outlook_rows(), start=4):
        ws.write(i, 0, CODE[b]); ws.write_number(i, 1, pk, num); ws.write_number(i, 2, e, num)
        ws.write_number(i, 3, f50, num); ws.write_number(i, 4, f90, num)
    ws.set_column(0, 0, 14); ws.set_column(1, 4, 16)
    wb.close()


# ------------------------------------------------------------------------------ markdown
FIELD_NOTES = """# Field notes: supply planning extracts

Sabine Crest Energy, Load Planning. Kept with the spring extracts; last revised 12 April 2027 (Craig Stewart).

Book codes are ERCOT weather zone codes: COAST, EAST, FWEST, NORTH, NCENT, SCENT, SOUTH, WEST. Times are Central
Prevailing Time. An ESI ID is the 17-digit premise identifier assigned by the distribution utility.

## enrollment_extract_20270409.csv
One row per enrolment entered in the enrolment system, open or closed.

| Field | Meaning |
|---|---|
| enrollment_id | Enrolment system key. |
| esi_id | Premise the enrolment serves. |
| record_type | NEW opens an enrolment. AMEND is an amendment of an open enrolment: it carries the ESI ID of the enrolment it amends and supersedes it. |
| book | Weather-zone book the premise is settled in. |
| tdsp | Distribution utility. |
| plan_code | Retail plan on the enrolment. |
| deposit_class | Credit deposit class (A, B, C; W waived). |
| meter_type | IDR for an interval data recorder; AMS for a profiled advanced meter. |
| premise_age_band | Year band the premise was built in, from the service application. |
| max_demand_kw | Maximum demand of the premise on the enrolment, kW. |
| broker_code | Broker who brought the enrolment; blank if direct. |
| entered_on | Date the row was entered. |
| start_date, end_date | First and last day of service under the enrolment. end_date is blank while the enrolment is open. |

## premise_register_20270409.csv
One row per ESI ID that has appeared on an enrolment. load_profile is the ERCOT load profile ID the
distribution utility assigns. naics_code is the customer's NAICS code from the service application. premise_status
is ACTIVE while an enrolment is open.

## switch_confirms_spring_2027.csv
Completed ERCOT market transactions that started service with Sabine Crest between 1 October 2026 and the extract:
switches (SWI) and move-ins (MVI). One row per completed transaction.

## idr_hourly_reads_summers_2017_2026.parquet
Interval reads for IDR premises while enrolled with Sabine Crest: weekdays June to September, hours ending 11 to 20.
kwh is the energy in the hour, which is also the average kW in the hour. hour_ending 17 is 16:00 to 17:00.

## zone_settled_load_s17_s26.csv
The book's load in each weather zone as settled by ERCOT at true-up, MWh per hour, every hour of June to September,
2017 to 2026. One column per book.

## ercot_summer_system_peaks.csv
ERCOT's settled summer system peak for each summer: date, hour ending and ERCOT load.

## billing_accounts.csv
Customer accounts on IDR summary billing and the ESI IDs billed under each, with the dates each premise was billed
under the account.

## bsaver_credits_2017_2026.csv
Business Saver credits as billed: one row per account and called window. credited_kwh and credit_usd are the
account's totals for the window across its enrolled sites; bill_month is the bill the credit posted to.

## Trading files
option_quotes_s27.csv: quotes received from brokers on the desk's quote line, one row per quote revision; premium in
the broker's convention. quote_decisions.csv: the desk's decision on each quote revision (ACCEPTED, DECLINED,
SUPERSEDED). desk_counterparties.csv: counterparties with the dates of their approval. book_zone_map.csv: ERCOT load zone each book is hedged and bought in, with effective dates.
trade_blotter_s27.csv: one row per booked trade amendment (amendment 0 is the original). confirm_match_log.csv:
matching status changes for each trade amendment. portfolio_books.csv: portfolio to book. trading_calendar_2027.csv:
calendar days with NERC holidays.
"""


def field_notes(path):
    with open(path, "w", newline="\n") as f:
        f.write(FIELD_NOTES)


def extract_log(path, entries):
    lines = ["# Extract log: summer 2027 block", "",
             "Supply Portfolio, Sabine Crest Energy. Files gathered for the 16 April 2027 Risk Committee. Each line: "
             "file, where it came from, what it covers.", ""]
    for f, txt in entries:
        lines.append(f"- `{f}`: {txt}")
    lines.append("")
    lines.append("Gathered by Supply Portfolio, 12 April 2027.")
    with open(path, "w", newline="\n") as fh:
        fh.write("\n".join(lines) + "\n")
