"""Read the shipped pack back into the frames analysis.py works on. The goldens are computed from these, never from
the world, so every golden figure is a figure the pack reproduces."""
from __future__ import annotations

import datetime as dt
import os
import re

import openpyxl
import pandas as pd

import common as K


def _csv(target, name):
    return pd.read_csv(os.path.join(target, name), dtype=str, keep_default_na=False)


def inventory(target) -> pd.DataFrame:
    inv = _csv(target, K.INVENTORY)
    out = pd.DataFrame({"sl_id": inv.sl_id, "zone": inv.zone, "segment": inv.main_segment,
                        "year_built": inv.year_built.astype(int), "record": inv.tap_card_material,
                        "cls": inv.sl_class, "basis": inv.class_basis, "replaced_on": inv.replaced_on})
    return out


def verifications(target) -> pd.DataFrame:
    """One row per verification: sl_id (CIS premise numbers through the crosswalk, padding restored), date, result
    class through the code list."""
    v = _csv(target, K.VERIFICATIONS)
    xw = _csv(target, K.CROSSWALK)
    codes = _csv(target, K.CODES)
    by_premise = dict(zip(xw.legacy_premise_no, xw.sl_id))
    cls = dict(zip(codes.result_code, codes.inventory_class))
    sl = [s if s else by_premise[p.zfill(8)] for s, p in zip(v.sl_id, v.premise_no)]
    date = [dt.datetime.strptime(d, "%m/%d/%Y").date() if src == "CIS" else dt.date.fromisoformat(d)
            for d, src in zip(v.verified_date, v.source)]
    return pd.DataFrame({"sl_id": sl, "date": date, "result": v.result_code.map(cls), "source": v.source,
                         "crew": v.crew, "method": v.method})


def unit_costs(target):
    wb = openpyxl.load_workbook(os.path.join(target, K.UNIT_COSTS), data_only=True)
    ws = wb.active
    prices, mob = {}, None
    for row in ws.iter_rows(values_only=True):
        if row[0] and re.fullmatch(r"RZ-\d\d", str(row[0])):
            prices[row[0]] = (float(row[2]), float(row[3]))
        elif row[0] and str(row[0]).startswith("Mobilisation"):
            mob = float(row[2])
    return prices, mob


def appropriation(target) -> float:
    import docx
    d = docx.Document(os.path.join(target, K.CAPITAL))
    for t in d.tables:
        for r in t.rows:
            cells = [c.text for c in r.cells]
            if cells[0].startswith("Lead service line replacement, construction"):
                return float(cells[1]) * 1e6
    raise KeyError("construction line")


def completions(target) -> pd.DataFrame:
    c = _csv(target, K.COMPLETIONS)
    for k in ("paid_base", "paid_extra_ft", "paid_restoration"):
        c[k] = c[k].astype(float)
    c["co_amount"] = c.co_amount.map(lambda x: float(x) if x else 0.0)
    utc = pd.to_datetime(c.completed_utc, utc=True)
    c["local_date"] = utc.dt.tz_convert("America/New_York").dt.date
    c["utc_date"] = utc.dt.date
    return c


def change_orders(target) -> pd.DataFrame:
    wb = openpyxl.load_workbook(os.path.join(target, K.CHANGE_ORDERS), data_only=True)
    ws = wb.active
    rows, head = [], None
    for row in ws.iter_rows(values_only=True):
        if row[0] == "CO no.":
            head = row
            continue
        if head and row[0]:
            rows.append(dict(zip(head, row)))
    df = pd.DataFrame(rows)
    df["sites"] = df["Sites covered"].map(lambda s: [x.strip() for x in s.split(",")])
    return df


def agreements(target) -> pd.DataFrame:
    a = _csv(target, K.AGREEMENTS)
    xw = _csv(target, K.CROSSWALK)
    by_premise = dict(zip(xw.legacy_premise_no, xw.sl_id))
    a["line"] = [s if s else by_premise[p.zfill(8)] for s, p in zip(a.sl_id, a.premise_no)]
    return a


def outreach(target) -> pd.DataFrame:
    return _csv(target, K.OUTREACH)
