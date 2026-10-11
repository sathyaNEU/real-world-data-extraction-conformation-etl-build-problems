"""Bellbird's documents: the ECL method, the scorecard migration note, the scenario workbook, the extract dictionary."""
import datetime as dt
import os
import shutil
import tempfile

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

import officex as OX
import world as W

SCEN_SHEET = "Q1 FY27 scenarios"        # the quarter to 30 September 2026, first quarter of the year to 30 June 2027
PRIOR_SHEET = "Q4 FY26 (prior)"


def methodology_docx(path):
    P = lambda t: ("p", t)
    blocks = [
        ("title", "CRM-ECL-04 Expected credit loss: unsecured personal loans (collective model)"),
        ("small", "Credit Risk Management policy and method library"),
        ("table", [["Document control", ""],
                   ["Document number", "CRM-ECL-04"],
                   ["Version", "4.1"],
                   ["Status", "Approved"],
                   ["Effective", "Reporting dates from 30 June 2026"],
                   ["Owner", "Head of Credit Risk Modelling"],
                   ["Approver", "Audit and Risk Committee"],
                   ["Next review", "May 2027"],
                   ["Related documents", "CRM-ECL-06 Stage 3 assessment; CRM-COL-02 Collections and write-off; "
                                         "HP-01 Financial hardship"]], [2600, 6400], 19),
        ("h1", "1. Scope and definitions"),
        P("1.1 This method sets the collective expected credit loss (ECL) allowance under AASB 9 Financial Instruments "
          "for PL36 and PL60 unsecured personal loans in stages 1 and 2. Loans in stage 3 (90 or more days past due) are "
          "assessed under CRM-ECL-06 and are outside this method."),
        P("1.2 A loan is in default when it is 90 or more days past due. The 90-day presumption in AASB 9 B5.5.37 is not "
          "rebutted for this portfolio, and loans in default are in stage 3."),
        P("1.3 Days past due are taken from the loan system at close of business on the reporting date."),
        P("1.4 Loans are written off under CRM-COL-02 in the month-end batch once 180 or more days past due. A loan that "
          "has been written off is closed in the loan system and carries no allowance."),
        P("1.5 The gross carrying amount of a loan is its current balance in the loan system."),
        ("h1", "2. Staging"),
        P("2.1 A loan is in stage 2 if, at the reporting date, any of the following holds: it is 30 or more days past due; "
          "a hardship arrangement is in force; or its current behaviour score is 60 or more points below its score at "
          "origination, both on the current (v4) scorecard. Otherwise it is in stage 1."),
        P("2.2 A hardship arrangement is in force if the reporting date falls on or between its start and end dates and it "
          "was not cancelled."),
        P("2.3 Origination scores recorded on the v3 scorecard are converted to v4 with the published migration formula "
          "before the comparison in 2.1."),
        P("2.4 Transfers between stages are symmetrical. A loan returns to stage 1 at the first reporting date on which "
          "none of the conditions in 2.1 holds, and a loan that is no longer 90 or more days past due is staged under 2.1. "
          "No probation period applies: the 2025 validation (MV-2025-07) found that one would move few loans and little "
          "allowance in this portfolio."),
        P("2.5 The 30 days past due backstop in AASB 9 5.5.11 is not rebutted."),
        ("h1", "3. Expected credit loss"),
        P("3.1 ECL = sum over future months k of PD(k) x LGD x EAD(k) x DF(k). Stage 1 loans use the next 12 months or the "
          "remaining term if shorter; stage 2 loans use the remaining term."),
        P("3.2 PD(k), the probability of default in the k-th month after the reporting date, comes from the model team's "
          "cumulative default curve for the loan's product. Stage 2 loans carry the significant-increase scalar of 2.4 on "
          "PD(k). A loan's remaining term is its contractual term less the number of monthly repayments that have fallen "
          "due on or before the reporting date."),
        P("3.3 LGD is 70 per cent for PL36 and 76 per cent for PL60."),
        P("3.4 EAD(k) is the balance outstanding at the start of month k when the current balance is repaid over the "
          "remaining term in equal monthly instalments at the contract rate."),
        P("3.5 DF(k) discounts month k at the contract rate: (1 + contract rate / 12) to the power -k."),
        P("3.6 ECL is calculated loan by loan and summed without rounding. Amounts are rounded only for reporting."),
        ("h1", "4. Forward-looking information"),
        P("Each scenario in the quarter's approved scenario workbook scales PD(k) by its multiplier. The allowance is the "
          "probability-weighted sum of the scenario ECLs. Group Finance sets the scenarios and their weights each quarter "
          "from the economic outlook, and the Audit and Risk Committee approves them before the reporting date."),
        ("h1", "5. Reporting"),
        P("The allowance is reported to the Audit and Risk Committee in Australian dollars by stage and product, with the "
          "coverage ratio (allowance over gross carrying amount, the sum of current balances) for each."),
        ("h1", "6. Post-model adjustments"),
        P("An overlay is made only for a risk this method does not capture. It is approved by the Audit and Risk Committee "
          "with the condition for its release and reported separately from the modelled allowance. The COVID-19 overlay was "
          "released at 30 June 2023 and no overlay has been approved since."),
        ("h1", "7. Limitations"),
        P("7.1 PD is set by product only. Differences in risk between grades, channels and states reach the allowance "
          "through staging, not through PD."),
        P("7.2 The stage 2 scalar was calibrated on arrears outcomes from 2023 to 2025 and is reviewed each year with the "
          "validation."),
        P("7.3 LGD is a portfolio average for each product and does not vary with time in arrears or with the size of the "
          "balance."),
        P("7.4 EAD assumes the current balance is repaid on schedule. Extra repayments, payouts and hardship variations "
          "after the reporting date are not modelled."),
        P("7.5 Scenario multipliers scale PD only. LGD is not stressed."),
        ("h1", "8. Governance"),
        P("8.1 Credit Risk Modelling reviews this method each year, and sooner if the products, the portfolio or the source "
          "data change materially. Model Validation reviews it independently at least every two years."),
        ("h2", "8.2 Sign-off"),
        ("table", [["Role", "Name and position", "Date"],
                   ["Prepared", "Matthew Gardner, Head of Credit Risk Modelling", "12 May 2026"],
                   ["Reviewed", "Adam Gordon, Manager, Model Validation", "15 May 2026"],
                   ["Endorsed", "Renee Foster, Chief Risk Officer", "19 May 2026"],
                   ["Approved", "Audit and Risk Committee (Chair: Stephanie Manning)", "26 May 2026"]], [1700, 5300, 2000], 19),
        ("h2", "8.3 Version history"),
        ("table", [["Version", "Date", "Change"],
                   ["1.0", "Nov 2019", "First issue, at the launch of the personal loan portfolio"],
                   ["2.0", "Jun 2020", "COVID-19 payment deferrals not treated as a significant increase in credit risk on "
                                       "their own; COVID-19 overlay introduced"],
                   ["3.0", "Jun 2023", "Default curves rebuilt on Bellbird's own loan history; COVID-19 overlay released"],
                   ["3.1", "Aug 2024", "Score condition moved to scorecard v4, with v3 origination scores converted"],
                   ["4.0", "Jun 2025", "Stage 2 scalar of 2.4 on PD(k) introduced (MV-2025-07)"],
                   ["4.1", "May 2026", "Hardship arrangements added to the stage 2 conditions; LGD for PL60 raised from 74 to "
                                       "76 per cent"]], [1100, 1300, 6600], 19),
    ]
    props = dict(title="CRM-ECL-04 Expected credit loss: unsecured personal loans", creator="Matthew Gardner",
                 modified_by="Matthew Gardner", created=dt.datetime(2019, 11, 4, 10, 21), modified=dt.datetime(2026, 5, 26, 13, 58),
                 revision=41, total_time=1174, company="Bellbird Personal Finance Ltd")
    OX.docx(path, blocks, ("Bellbird Personal Finance Ltd", "CRM-ECL-04 Expected credit loss: unsecured personal loans"),
            ("Version 4.1  |  Internal  |  Uncontrolled when printed", "Page {PAGE} of {NUMPAGES}"), props, seed=4041, sessions=14)


def scorecard_pdf(path):
    blocks = [
        ("title", "Behaviour and application scorecard v4: migration note"),
        ("small", "Credit Risk Modelling  |  20 June 2024"),
        ("p", "From 1 July 2024 all application and behaviour scores are produced on scorecard v4. Scores recorded before "
              "that date are on v3 and are not restated in the loan system."),
        ("p", "v4 is a recalibration of v3 on the same characteristics with a wider points scale. A v3 score is put on the v4 "
              "scale as v4 = 1.12 x v3 - 18, rounded to the nearest whole point."),
        ("p", "Worked example: a v3 origination score of 615 converts to 1.12 x 615 - 18 = 670.8, which is 671 on v4."),
        ("p", "Application grade bands were re-set on v4 to hold the approval rate, as below. The decline rules and the "
              "overrides that use them are in Credit Policy CP-04."),
        ("table", [["Grade", "Minimum score, v3", "Minimum score, v4"],
                   ["A", "700", "766"], ["B", "655", "716"], ["C", "615", "671"], ["D", "580", "632"], ["E", "545", "592"]],
         [1800, 2600, 2600], 19),
        ("p", "Validation of v4 on the 2022 to 2023 applications (MV-2024-03) found the rank ordering of the two versions "
              "equivalent, with a Gini coefficient of 0.61 for v4 against 0.59 for v3."),
    ]
    props = dict(title="Behaviour and application scorecard v4: migration note", creator="Troy Olson",
                 modified_by="Troy Olson", created=dt.datetime(2024, 6, 17, 11, 6), modified=dt.datetime(2024, 6, 20, 15, 2),
                 revision=5, total_time=96, company="Bellbird Personal Finance Ltd")
    tmp = tempfile.mkdtemp()
    src = os.path.join(tmp, "note.docx")
    OX.docx(src, blocks, ("Bellbird Personal Finance", "Credit Risk Modelling"),
            ("Scorecard v4 migration note  |  June 2024", "Page {PAGE} of {NUMPAGES}"), props, font="Arial", seed=2406,
            sessions=4, margins=(1418, 1418, 1247, 1247))
    OX.pdf_from_docx(src, path)
    shutil.rmtree(tmp)


def scenarios_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SCEN_SHEET
    ws.append(["IFRS 9 / AASB 9 forward-looking scenarios, reporting date 30 September 2026"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Approved by the Audit and Risk Committee, 18 September 2026"])
    ws.append([])
    hdr = ["Scenario", "Probability weight", "PD multiplier", "Unemployment rate, peak (%)", "Cash rate, end FY27 (%)",
           "GDP growth, FY27 (%)"]
    ws.append(hdr)
    for j in range(1, len(hdr) + 1):
        ws.cell(row=4, column=j).font = Font(bold=True)
        ws.cell(row=4, column=j).fill = PatternFill("solid", fgColor="E5E9F0")
        ws.cell(row=4, column=j).alignment = Alignment(wrap_text=True, vertical="top")
    macro = {"Base": (4.4, 3.35, 2.1), "Downside": (6.1, 2.60, 0.4), "Upside": (3.9, 3.60, 2.8)}
    for name, w, m in W.SCENARIOS:
        ws.append([name, w, m] + list(macro[name]))
    for r in range(5, 8):
        ws.cell(row=r, column=2).number_format = "0%"
        ws.cell(row=r, column=3).number_format = "0.00"
        for c in (4, 5, 6):
            ws.cell(row=r, column=c).number_format = "0.0" if c != 5 else "0.00"
    ws.append([])
    ws.append(["Source: Group Finance economic outlook, September 2026. Multipliers apply to PD only (CRM-ECL-04 section 4)."])
    ws["A9"].font = Font(italic=True, color="5B6670")
    for col, wd in zip("ABCDEF", [14, 13, 11, 16, 15, 13]):
        ws.column_dimensions[col].width = wd
    ws.row_dimensions[4].height = 30
    ws2 = wb.create_sheet(PRIOR_SHEET)
    ws2.append(["Scenarios used at 30 June 2026"])
    ws2["A1"].font = Font(bold=True, size=12)
    ws2.append(["Approved by the Audit and Risk Committee, 18 June 2026"])
    ws2.append([])
    ws2.append(["Scenario", "Probability weight", "PD multiplier"])
    for j in range(1, 4):
        ws2.cell(row=4, column=j).font = Font(bold=True)
        ws2.cell(row=4, column=j).fill = PatternFill("solid", fgColor="E5E9F0")
    for name, w, m in [("Base", 0.60, 1.00), ("Downside", 0.25, 1.25), ("Upside", 0.15, 0.85)]:
        ws2.append([name, w, m])
    for r in range(5, 8):
        ws2.cell(row=r, column=2).number_format = "0%"
        ws2.cell(row=r, column=3).number_format = "0.00"
    for col, wd in zip("ABC", [14, 13, 11]):
        ws2.column_dimensions[col].width = wd
    OX.xlsx_from_openpyxl(wb, path, "Tamara Farrell", "Tamara Farrell", dt.datetime(2026, 9, 8, 9, 41, 17),
                          dt.datetime(2026, 9, 18, 16, 52, 38), seed=918)


def dictionary_txt(path):
    txt = """BELLBIRD PERSONAL FINANCE LTD
QUARTER-END ECL RUN, 30 SEPTEMBER 2026: LOAN SYSTEM EXTRACTS
Produced by job ECL_QE_BOOK on 1 October 2026 at 07:30, after the September month-end batch.
Credit Risk Reporting keeps this dictionary; changes to an extract are recorded against the job.


1. personal_loan_book_2026-09-30.csv
Every loan open in the loan system at close of business on 30 September 2026, one row per loan, including loans in
arrears. Loans are written off in the month-end batch once 180 or more days past due, so loans written off in the
September batch are not in the extract. Sorted by loan number.

  loan_no                  loan number, allocated at application: BPL, six digits and a check digit
  product                  PL36 (36-month term) or PL60 (60-month term)
  risk_grade               grade at approval, A (lowest risk) to E, from the application score bands in force on the day
  drawdown_date            date the loan was funded
  first_repayment_date     date of the first scheduled repayment; later repayments are scheduled on the same day of each
                           month. A repayment scheduled on a weekend or NSW public holiday falls due on the next
                           business day
  term_months              contractual term in months
  contract_rate            fixed annual contract rate, decimal
  original_amount          amount financed, whole dollars, including the establishment fee
  current_balance          account balance at close of business: principal plus interest and fees charged, less
                           repayments received (AUD). Interest accrues daily and is charged on each scheduled
                           repayment date
  days_past_due            days since the due date of the oldest repayment not paid in full; 0 when up to date
  origination_score        application score at origination, on the scorecard named in origination_scorecard
  origination_scorecard    v3 (applications before 1 July 2024) or v4
  current_score_v4         behaviour score at the reporting date, scorecard v4
  state                    state of the borrower's residential address
  channel                  origination channel: Broker, Online or Direct


2. hardship_register_2026-09-30.xlsx, sheet Arrangements
Every hardship arrangement recorded since the loan system went live in October 2019, one row per arrangement,
including arrangements on loans since repaid or written off. A loan can have more than one arrangement. Case numbers
are allocated when a request is received; requests declined at assessment keep their number and are not listed.
Arrears at the start of an arrangement are capitalised.

  case_ref                 hardship case number
  loan_no                  loan number
  received_date            date the request was received
  hardship_reason          reason recorded at assessment
  arrangement_type         Payment pause (no repayments for the period) or Reduced repayments
  arrangement_start        first day of the varied repayments
  arrangement_end          last day of the varied repayments, as approved
  agreed_repayment_aud     monthly repayment agreed under a reduced repayment arrangement; 0 for a pause
  arrangement_status       Approved (not yet started), Active, Completed or Cancelled
  cancelled_date           date the arrangement was cancelled, at the customer's request or after varied repayments
                           were missed; a cancelled arrangement keeps its approved dates


3. write_offs_2025-07_2026-09.csv
Loans written off in each month-end batch, July 2025 to September 2026, and recoveries received.

  month                    month of the batch (YYYY-MM)
  loans_written_off        number of loans written off
  gross_write_off_aud      balances written off (AUD)
  recoveries_aud           amounts received in the month on loans written off in any earlier month (AUD)


Other files in the run folder (not loan system extracts)
  ifrs9_scenarios_2026Q3.xlsx            Group Finance: scenario weights and PD multipliers approved by the Audit and
                                         Risk Committee, one sheet per quarter
  pd_curve_personal_loans_2026-06.csv    Credit Risk Modelling: default curves; described in the file's header lines
  pd_model_development_sample.parquet    Credit Risk Modelling: the curves' development data; described in the file's
                                         metadata
  arrears_roll_rates_fy26.xlsx           Collections MI: month-end roll rates between arrears buckets, FY26
"""
    with open(path, "w", newline="\r\n") as f:
        f.write(txt)
