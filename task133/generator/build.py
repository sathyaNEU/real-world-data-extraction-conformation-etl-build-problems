"""Build task133's evidence pack into <out>/ (default: task133/target)."""
import argparse
import datetime as dt
import json
import os
import re
import shutil
import sys

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import world as W
import sim as SM
import extracts as X
import docs as D
import distract as Q
import officex as OX
import postprocess as P

TASK = os.path.dirname(HERE)
FILES = {
    "book": "personal_loan_book_2026-09-30.csv",
    "june": "personal_loan_book_2026-06-30.csv",
    "curve": "pd_curve_personal_loans_2026-06.csv",
    "method": "CRM-ECL-04_expected_credit_loss_personal_loans.docx",
    "scen": "ifrs9_scenarios_2026Q3.xlsx",
    "hardship": "hardship_register_2026-09-30.xlsx",
    "note": "scorecard_v4_migration_note.pdf",
    "dictionary": "ecl_extract_dictionary.txt",
    "dev": "pd_model_development_sample.parquet",
    "scores": "behaviour_score_runs_2026-03_2026-08.parquet",
    "defaults": "collections_defaults_2023-24.csv",
    "receipts": "collections_receipts_2023-07_2026-09.csv",
    "sale": "debt_sale_2026-03_schedule.xlsx",
    "writeoffs": "write_offs_2025-07_2026-09.csv",
    "rolls": "arrears_roll_rates_fy26.xlsx",
}
DISTRACTORS = ["writeoffs", "rolls"]
# Sydney wall-clock times each file was last saved
MTIMES = {
    "book": (2026, 10, 1, 7, 34), "june": (2026, 7, 1, 7, 33), "curve": (2026, 6, 24, 16, 2),
    "method": (2026, 8, 26, 9, 14), "scen": (2026, 9, 18, 16, 52), "hardship": (2026, 10, 1, 8, 2),
    "note": (2024, 6, 20, 15, 5), "dictionary": (2026, 10, 2, 10, 6), "dev": (2026, 6, 2, 11, 45),
    "scores": (2026, 9, 2, 10, 17), "defaults": (2026, 10, 2, 9, 41), "receipts": (2026, 10, 2, 9, 38),
    "sale": (2026, 3, 27, 14, 5), "writeoffs": (2026, 10, 1, 8, 5), "rolls": (2026, 7, 14, 9, 50),
}
SCORE_RUN_ENDS = [SM.dn(dt.date(2026, m + 1, 1)) - 1 for m in range(3, 9)]     # the March to August 2026 month ends
CURVE_HEADER = """# Bellbird Personal Finance - Credit Risk Modelling
# PLPD-2026.06 personal loan default curves, PL36 and PL60; approved by the Model Risk Committee on 24 June 2026
# Developed on the PLPD-2026.06 development sample, performance to 31 May 2026: every PL36 loan drawn October 2019 to March 2023 and every PL60 loan drawn October 2019 to March 2021, the loan system's go-live to the last complete-term vintages
# Default: 90 or more days past due at a month end. All risk grades pooled
# month_on_book: number of monthly repayments fallen due; repayment dates during a payment pause count
# cumulative_pd: share of the loans drawn in the development window that had defaulted by the end of that month on book
"""
DEV_META = {
    "dataset": "PLPD-2026.06 development sample",
    "owner": "Credit Risk Modelling",
    "extracted": "2026-06-02",
    "population": "every PL36 loan drawn October 2019 to March 2023 and every PL60 loan drawn October 2019 to March 2021: "
                  "the loan system's go-live to the last complete-term vintages",
    "performance_to": "2026-05-31",
    "outcome": "default (90 or more days past due at a month end), prepaid (repaid in full before the end of the term) or "
               "matured",
    "outcome_month_on_book": "monthly repayments fallen due at default, payout or maturity, repayment dates during a "
                             "payment pause included",
}
SCORE_META = {
    "dataset": "Behaviour score v4: month-end scoring runs",
    "owner": "Behaviour Scoring",
    "extracted": "2026-09-02",
    "runs": "the March, April, May, June, July and August 2026 month ends",
    "run_date": "the day the run was made: each month-end run is made early on the first day of the following month and "
                "scores every open loan on its position at close of business on the month end",
    "population": "every loan open at the month end, from its first repayment date",
    "score_v4": "behaviour score, scorecard v4",
}


def register_xlsx(hd, path):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Arrangements"
    cols = list(hd.columns)
    ws.append(cols)
    for j in range(1, len(cols) + 1):
        c = ws.cell(row=1, column=j)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
        c.alignment = Alignment(vertical="top")
    for rec in hd.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and v != v) else v for v in rec])
    n = len(hd) + 1
    for r in range(2, n + 1):
        for j, name in enumerate(cols, 1):
            if name.endswith("_date") or name in ("arrangement_start", "arrangement_end"):
                ws.cell(row=r, column=j).number_format = "dd/mm/yyyy"
            elif name.endswith("_aud"):
                ws.cell(row=r, column=j).number_format = "#,##0.00"
    for col, w in zip("ABCDEFGHIJ", [11, 12, 13, 21, 19, 16, 16, 20, 17, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (OX.col_letter(len(cols)), n)
    # written by the job under its service account; an analyst in Credit Risk Reporting added the header fill, the
    # frozen header row and the filter before the run folder was released
    OX.xlsx_from_openpyxl(wb, path, W.JOB_ACCOUNT, "Kayla Stokes", dt.datetime(2026, 10, 1, 7, 36, 5),
                          dt.datetime(2026, 10, 1, 8, 2, 44), seed=1001)


def curve_csv(curves, path):
    lines = [CURVE_HEADER.rstrip("\n"), "product,month_on_book,cumulative_pd,model_version"]
    for prod, term in W.PRODUCTS.items():
        for m, v in enumerate(curves[prod]):
            lines.append("%s,%d,%.6f,PLPD-2026.06" % (prod, m, v))
    with open(path, "w") as f:
        f.write("\n".join(lines) + "\n")


def parquet(df, path, meta):
    import pyarrow as pa
    import pyarrow.parquet as pq
    t = pa.Table.from_pandas(df, preserve_index=False)
    t = t.replace_schema_metadata({k: v for k, v in meta.items()})
    pq.write_table(t, path)
    b = open(path, "rb").read()
    b = re.sub(rb"parquet-cpp-arrow version \d{2}\.\d\.\d", b"parquet-cpp-arrow version 16.1.0", b)
    open(path, "wb").write(b)


def score_frame(S, runs):
    import pandas as pd
    parts = []
    for t in sorted(runs):
        rows, sc = runs[t]
        parts.append(pd.DataFrame(dict(loan_no=S.L["loan_no"][rows], run_date=SM.dd(t + 1), score_v4=sc.astype(np.int32))))
    out = pd.concat(parts, ignore_index=True)
    return out.sort_values(["run_date", "loan_no"], kind="stable").reset_index(drop=True)


def loss_history(S):
    """the simulated loss history sits where the filed LGDs say it should: written-off cohorts with three years behind
    them have recovered roughly a quarter of the balance written off (agency plans, settlements, odd payments and, for
    the accounts sold in March 2026, their share of the sale), and about one default in six is resolved (settled, caught
    up or put on a collections arrangement) before the 180-day write-off"""
    C = X.collections(S)
    R = S.R
    wo, bal = C["wo"], C["wbal"]
    got = np.zeros(len(wo))
    pos = {int(i): k for k, i in enumerate(wo)}
    for i, d, a in zip(C["ledger"]["_i"], C["ledger"]["_d"], C["ledger"]["_amt"]):
        if i >= 0 and i in pos and d > S.close[i]:
            got[pos[i]] += a
    for k, (d, a, kind) in ((pos[i], f) for i, fl in C["flows_by"].items() for f in fl):
        if d < SM.dn(dt.date(2023, 7, 1)) and d > S.close[wo[k]]:
            got[k] += a                          # before the receipts extract begins
    price = W.SALE_PRICE_CENTS / 100.0
    for i, face in C["sold"].items():
        got[pos[i]] += face * price
    w = S.close[wo]
    mature = w <= SM.dn(dt.date(2023, 9, 30))
    rate = got[mature].sum() / bal[mature].sum()
    assert 0.20 <= rate <= 0.30, "mature write-off cohorts recover %.3f of the balance written off" % rate
    young = (w > SM.dn(dt.date(2025, 9, 30)))
    assert got[young].sum() / bal[young].sum() < 0.6 * rate, "a young cohort shows a partial picture"
    reached = (S.j0 > 0) & (S.e_def <= R - 200)
    resolved = reached & (S.close_type != 3)
    share = resolved.sum() / reached.sum()
    assert 0.12 <= share <= 0.25, "defaults resolved before write-off %.3f" % share
    return rate, share


def build(out):
    if os.path.exists(out):
        shutil.rmtree(out)
    os.makedirs(out)
    S = SM.simulate()
    p = lambda k: os.path.join(out, FILES[k])
    bk = X.book(S, np.random.default_rng(W.SEED + 1))
    hide = ["loan_key", "_bal", "_dpd"]
    bk.drop(columns=hide).to_csv(p("book"), index=False, lineterminator="\n")
    jb = X.frame_at(S, SM.dn(W.PRIOR_QE))
    jb.drop(columns=hide).to_csv(p("june"), index=False, lineterminator="\n")
    parquet(score_frame(S, X.score_runs(S, SCORE_RUN_ENDS)), p("scores"), SCORE_META)
    hd = X.hardship(S, np.random.default_rng(W.SEED + 2))
    register_xlsx(hd, p("hardship"))
    dev = X.development_sample(S)
    curve_csv(X.curve(dev), p("curve"))
    parquet(dev, p("dev"), DEV_META)
    D.methodology_docx(p("method"))
    D.scenarios_xlsx(p("scen"))
    D.scorecard_pdf(p("note"))
    P.normalize_pdf(p("note"), dt.datetime(*MTIMES["note"]))
    D.dictionary_txt(p("dictionary"))
    C = X.collections(S)
    X.fy24_defaults(S).drop(columns=["_i", "_dme"]).to_csv(p("defaults"), index=False, lineterminator="\n")
    C["ledger"].drop(columns=["_i", "_d", "_amt"]).to_csv(p("receipts"), index=False, lineterminator="\n")
    D.sale_xlsx(X.sale_schedule(S), C["consideration"], p("sale"))
    X.write_offs(S).to_csv(p("writeoffs"), index=False, lineterminator="\n")
    loss_history(S)
    Q.roll_rates_xlsx(p("rolls"), X.roll_rates(S))
    for k in FILES:
        ts = W.posix(dt.datetime(*MTIMES[k]))
        os.utime(p(k), (ts, ts))
    return S, bk, hd


def write_metadata():
    meta = {
        "task": "task133",
        "title": "Bellbird personal loans: the 30 September 2026 ECL allowance",
        "domain": "Banking, Credit & Financial Risk",
        "subdomain": "expected-credit-loss-provisioning",
        "objective": "Forecasting & Predictive Modeling",
        "as_of": "2026-09-30",
        "deliverables": ["ecl_allowance_2026Q3.xlsx", "board_ecl_q3_2026.pptx"],
        "input_files": [{"path": FILES[k], "format": FILES[k].rsplit(".", 1)[1]} for k in sorted(FILES, key=lambda k: FILES[k])],
        "distractors": [FILES[k] for k in DISTRACTORS],
        "provenance": "Fully constructed pack (dataset-generation section 2): every extract is generated by task133/generator "
                      "from a simulated personal loan book for a fictional Australian non-bank lender; documents written "
                      "for the fiction. No real lender's or borrower's data.",
        "license": "Constructed data, released with the task (CC BY 4.0).",
        "created": "2026-10-11",
    }
    with open(os.path.join(TASK, "metadata.json"), "w") as f:
        json.dump(meta, f, indent=2)
        f.write("\n")


ZIP_TIME = (2026, 10, 2, 10, 30, 0)


def zip_target(src, path):
    """the bundle as shipped: every file flat, in name order, one fixed timestamp, so a rebuild is byte-identical"""
    import zipfile
    with zipfile.ZipFile(path, "w") as z:
        for name in sorted(os.listdir(src)):
            zi = zipfile.ZipInfo(name, date_time=ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.create_system = 3
            zi.external_attr = 0o644 << 16
            with open(os.path.join(src, name), "rb") as f:
                z.writestr(zi, f.read())


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(TASK, "target"))
    ap.add_argument("--zip", default=None, help="where to write the zipped bundle (default: target.zip beside target/)")
    a = ap.parse_args()
    build(a.out)
    if a.out == os.path.join(TASK, "target"):
        write_metadata()
        zip_target(a.out, os.path.join(TASK, "target.zip"))
    elif a.zip:
        zip_target(a.out, a.zip)
    print("built", a.out)
