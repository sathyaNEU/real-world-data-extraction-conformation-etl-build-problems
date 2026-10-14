"""Independent check of every graded figure, every rung of the main ladder and every ask stop, from the shipped bytes
only: numpy difference arrays for the record's default and repayment rates and a loan by month grid for the allowance,
plain dictionaries and the csv module for the two asks, its own date arithmetic and its own scorecard conversion read from
the migration note. Nothing here imports the generator."""
import csv
import datetime as dt
import os
import re
import sys
from collections import defaultdict

import numpy as np
import openpyxl
import pandas as pd
from pypdf import PdfReader

TASK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
T = os.path.join(TASK, "target")
G = os.path.join(TASK, "golden")
REP = dt.date(2026, 9, 30)
JUNE = dt.date(2026, 6, 30)
TOP = 81                                     # rates are read to month on book 81, past every term on the book


def mi(s):
    """month index of a date or a YYYY-MM string"""
    s = str(s)
    return int(s[:4]) * 12 + int(s[5:7]) - 1


def mi_or(v):
    """month index, or -1 for an empty field"""
    return -1 if v is None or (isinstance(v, float) and v != v) or v is pd.NaT else mi(v)


RM = mi(REP)
GO = mi("2019-10")                           # the loan system went live on 1 October 2019


def read_inputs(target):
    bk = pd.read_csv(os.path.join(target, "personal_loan_book_2026-09-30.csv"), parse_dates=["first_repayment_date", "drawdown_date"])
    hd = pd.read_excel(os.path.join(target, "hardship_register_2026-09-30.xlsx"),
                       parse_dates=["arrangement_start", "arrangement_end", "cancelled_date"])
    sc = openpyxl.load_workbook(os.path.join(target, "ifrs9_scenarios_2026Q3.xlsx"), data_only=True)["Q1 FY27 scenarios"]
    mult = 0.0
    for r in sc.iter_rows(min_row=5, values_only=True):
        if not r or r[0] is None:
            break
        mult += float(r[1]) * float(r[2])
    note = " ".join(p.extract_text() for p in PdfReader(os.path.join(target, "scorecard_v4_migration_note.pdf")).pages)
    a, b = map(float, re.search(r"v4 = ([\d.]+) x v3 - ([\d.]+)", note).groups())
    return bk, hd, mult, (a, b)


def in_force(hd, day):
    day = pd.Timestamp(day)
    sel = (hd["arrangement_start"] <= day) & (hd["arrangement_end"] >= day) & (hd["cancelled_date"].isna() | (hd["cancelled_date"] > day))
    return set(hd.loc[sel, "loan_no"])


def stages(frame, hd, conv, day, score):
    a, b = conv
    o4 = np.where(frame["origination_scorecard"] == "v3", np.floor(a * frame["origination_score"] - b + 0.5), frame["origination_score"])
    fall = o4 - np.asarray(score, float)
    hard = frame["loan_no"].isin(in_force(hd, day)).to_numpy()
    st = np.where((frame["days_past_due"].to_numpy() >= 30) | hard | (np.nan_to_num(fall, nan=-1e9) >= 60), 2, 1)
    return np.where(frame["days_past_due"].to_numpy() >= 90, 3, st)


def record_rates(target, entry="transfer", grade=False, clock="draw"):
    """for each product (and grade), the default rate and the repayment rate of every month on book m = 1..81: of the
    loans the record holds that are open and not in default at the start of m, the share first defaulting in m and the
    share closing in m without a default. entry 'transfer': a loan taken over at the go-live is held from its first month
    on book after September 2019; 'drawdown': every loan from month on book 1; 'own': the taken-over loans left out.
    clock 'first' counts months on book from the first repayment month"""
    rec = pd.read_parquet(os.path.join(target, "loan_record_2026-09-30.parquet"))
    draw = np.array([mi(x) for x in rec["drawdown_date"]])
    d0 = np.array([mi(x) for x in rec["first_repayment_date"]]) if clock == "first" else draw
    dflt = np.array([mi_or(x) for x in rec["default_month"]])
    shut = np.array([mi_or(x) for x in rec["closed_date"]])
    kind = np.where(dflt >= 0, 1, np.where(shut >= 0, 2, 0))
    ev_m = np.where(dflt >= 0, dflt, np.where(shut >= 0, shut, RM))
    held = np.maximum(0, GO - 1 - draw)      # months on book the earlier platform ran before the take-over
    assert ((ev_m - draw > held) | (kind == 0) | (held == 0)).all(), "every taken-over loan's default or closing falls after the take-over"
    x = ev_m - d0
    lo = held - (d0 - draw) if entry == "transfer" else np.full(len(rec), -10 ** 6)
    keep = held == 0 if entry == "own" else np.ones(len(rec), bool)
    keys = (rec["product"] + ("/" + rec["risk_grade"] if grade else "")).to_numpy()
    out = {}
    for k in sorted(set(keys)):
        sel = (keys == k) & keep
        xx, kk, ll = x[sel], kind[sel], lo[sel]
        diff = np.zeros(TOP + 3)
        np.add.at(diff, np.clip(ll + 1, 1, TOP + 2), 1)
        np.add.at(diff, np.clip(xx + 1, 1, TOP + 2), -1)
        risk = np.cumsum(diff)[:TOP + 1]
        ev = (xx >= 1) & (xx <= TOP)
        dd = np.bincount(xx[ev & (kk == 1)], minlength=TOP + 1)[:TOP + 1]
        cc = np.bincount(xx[ev & (kk == 2)], minlength=TOP + 1)[:TOP + 1]
        with np.errstate(invalid="ignore", divide="ignore"):
            d_ = np.where(risk > 0, dd / np.maximum(risk, 1), 0.0)
            c_ = np.where(risk > 0, cc / np.maximum(risk, 1), 0.0)
        d_[0] = c_[0] = 0.0
        out[k] = (d_, c_)
    return out


def crude_rates(target):
    """the audit's crude route: each product's cumulative default share by month on book over the loans in the record
    whose outcome is known (defaulted or closed)"""
    rec = pd.read_parquet(os.path.join(target, "loan_record_2026-09-30.parquet"))
    rec = rec[(rec["status"] == "Closed") | rec["default_month"].notna()]
    out = {}
    for p, g in rec.groupby("product"):
        d0 = np.array([mi(x) for x in g["drawdown_date"]])
        dm = np.array([mi_or(x) for x in g["default_month"]])
        x = (dm - d0)[dm >= 0]
        out[p] = np.cumsum(np.bincount(np.clip(x, 0, TOP), minlength=TOP + 1)) / len(g)
    return out


def allowance(target, how="gold"):
    """the collective allowance by stage. how: 'gold' (the record's rates, taken-over loans held from the go-live),
    'drawdown' (every loan counted from its drawdown), 'own' (the taken-over loans left out), 'km' (the default rate
    alone, repayment ignored), 'grade' (rates by product and grade, from drawdown), 'first' (months on book from the
    first repayment), 'crude' (the audit's crude conditional curve), 'r4b' (repayment counted again on EAD)"""
    bk, hd, mult, conv = read_inputs(target)
    rep = pd.Timestamp(REP)
    fr = bk["first_repayment_date"]
    due = (rep.year - fr.dt.year) * 12 + (rep.month - fr.dt.month) + (fr.dt.day <= rep.day).astype(int)
    due = due.clip(lower=0)
    rem = (bk["term_months"] - due).where(bk["term_months"] - due > 0, 1).to_numpy()
    base = bk["first_repayment_date"] if how == "first" else bk["drawdown_date"]
    a_now = RM - np.array([mi(x) for x in base.dt.strftime("%Y-%m")])
    st = stages(bk, hd, conv, REP, bk["current_score_v4"])
    lgd = bk["product"].map({"PL36": 0.70, "PL60": 0.76}).to_numpy()
    entry = {"gold": "transfer", "drawdown": "drawdown", "own": "own", "km": "transfer", "grade": "drawdown",
             "first": "transfer", "crude": None, "r4b": "transfer"}[how]
    rates = record_rates(target, entry, grade=how == "grade", clock="first" if how == "first" else "draw") if entry else None
    crude = crude_rates(target) if how == "crude" else None
    keys = (bk["product"] + ("/" + bk["risk_grade"] if how == "grade" else "")).to_numpy()
    ecl = np.zeros(len(bk))
    kk = np.arange(1, 61)[None, :]
    for key in sorted(set(keys)):
        idx = np.where((keys == key) & (st < 3))[0]
        if not len(idx):
            continue
        aa, rr = a_now[idx], rem[idx]
        H = np.where(st[idx] == 2, rr, np.minimum(12, rr))
        mask = kk <= H[:, None]
        m = np.minimum(aa[:, None] + kk, TOP)
        if how == "crude":
            c = crude[key]
            pdk = (c[m] - c[m - 1]) / (1 - c[np.minimum(aa, TOP)])[:, None]
        else:
            d, c = rates[key]
            exitm = d[m] + (0.0 if how == "km" else c[m])
            surv = np.cumprod(np.hstack([np.ones((len(idx), 1)), 1 - exitm[:, :-1]]), axis=1)
            pdk = d[m] * surv
        rate = bk["contract_rate"].to_numpy()[idx] / 12.0
        bal = bk["current_balance"].to_numpy()[idx]
        pmt = bal * rate / (1 - (1 + rate) ** (-rr.astype(float)))
        g = (1 + rate)[:, None] ** (kk - 1)
        ead = bal[:, None] * g - pmt[:, None] * (g - 1) / rate[:, None]
        if how == "r4b":                       # repayment counted again on EAD, on top of the competing exit
            ead = ead * np.cumprod(np.hstack([np.ones((len(idx), 1)), 1 - rates[key][1][m][:, :-1]]), axis=1)
        df = (1 + rate)[:, None] ** (-kk)
        scal = np.where(st[idx] == 2, 2.4, 1.0)[:, None]
        ecl[idx] = np.where(mask, pdk * scal * mult * lgd[idx][:, None] * ead * df, 0.0).sum(axis=1)
    out = {"total": ecl[st < 3].sum(), "stage1": ecl[st == 1].sum(), "stage2": ecl[st == 2].sum()}
    return out, st, bk


def _xl_rows(path):
    ws = openpyxl.load_workbook(path, read_only=True).worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = list(next(it))
    return [dict(zip(head, r)) for r in it]


def _day(v):
    return v.date() if hasattr(v, "date") else dt.date.fromisoformat(str(v)[:10])


def owners(target, led, refs="asof"):
    """the loan each receipt counts against: an agency remittance quotes Wirraway's reference and belongs to the
    placement that carried it on the day it was received ('asof'); 'posted' keeps the ledger's loan number"""
    if refs == "posted":
        return [r["loan_no"] for r in led]
    held = defaultdict(list)
    with open(os.path.join(target, "wirraway_placements_2026-09-30.csv"), newline="") as f:
        for r in csv.DictReader(f):
            p = dt.datetime.strptime(r["Placed"], "%d/%m/%Y").date()
            c = dt.datetime.strptime(r["Closed"], "%d/%m/%Y").date() if r["Closed"] else dt.date.max
            held[r["Agency Ref"]].append((p, c, r["Client Ref"]))
    out = []
    for r in led:
        if r["source"] != "Agency remittance":
            out.append(r["loan_no"])
            continue
        ref = re.search(r"ref (\d+)$", r["narrative"]).group(1)
        d = dt.date.fromisoformat(r["received_date"])
        hit = [l for p, c, l in held[ref] if p <= d <= c]
        assert len(hit) == 1, (ref, d)
        out.append(hit[0])
    return out


def not_cured(target, rows):
    """cases Collections closed as cured on a collections hardship arrangement that never started (received after the
    default month end, starting on the case's resolution date, cancelled before that start), and those among them whose
    arrears were never made good: the loan record closes them as written off or settled for less"""
    reg = _xl_rows(os.path.join(target, "hardship_register_2026-09-30.xlsx"))
    rec = pd.read_parquet(os.path.join(target, "loan_record_2026-09-30.parquet"), columns=["loan_no", "closure_reason"])
    gone = set(rec.loc[rec["closure_reason"].isin(["Written off", "Settled"]), "loan_no"])
    cured = {x["loan_no"]: x for x in rows if x["resolution"] == "Cured"}
    hit = set()
    for r in reg:
        x = cured.get(r["loan_no"])
        if x is None or r["arrangement_status"] != "Cancelled":
            continue
        y, m = int(x["default_month"][:4]), int(x["default_month"][5:7])
        dme = dt.date(y + (m == 12), m % 12 + 1, 1) - dt.timedelta(days=1)
        start, cancel = _day(r["arrangement_start"]), _day(r["cancelled_date"])
        if _day(r["received_date"]) > dme and start == dt.date.fromisoformat(x["resolution_date"]) and cancel < start:
            hit.add(x["loan_no"])
    return hit, hit & gone


def lgd(target, cures="gold", refs="asof", sale_in=True):
    """realised loss on the 2023-24 defaults by product and quarter of the default month, in per cent (CRM-ECL-04 3.3):
    the quarter's balance at default not recovered, summed over its defaults, over their total balance at default"""
    rows = list(csv.DictReader(open(os.path.join(target, "collections_defaults_2023-24.csv"))))
    led = list(csv.DictReader(open(os.path.join(target, "collections_receipts_2023-07_2026-09.csv"))))
    ws = openpyxl.load_workbook(os.path.join(target, "debt_sale_2026-03_schedule.xlsx"), data_only=True).worksheets[0]
    vals = list(ws.iter_rows(values_only=True))
    head = {r[0]: r[1] for r in vals if r[0] and r[0] != "loan_no" and not str(r[0]).startswith("BPL")}
    price = float(str(head["Price"]).split()[0]) / 100.0
    face = {r[0]: float(r[3]) for r in vals if r[0] and str(r[0]).startswith("BPL")}
    loans = {r["loan_no"]: r for r in rows}
    withdrawn, stale = not_cured(target, rows)
    # CRM-ECL-04 3.3: a loan the loan record shows written off or settled for less did not return to performing,
    # whatever resolution Collections recorded
    shut = pd.read_parquet(os.path.join(target, "loan_record_2026-09-30.parquet"), columns=["loan_no", "closure_reason"])
    gone = {l for l, c in zip(shut["loan_no"], shut["closure_reason"]) if c in ("Written off", "Settled")}
    got_pv = defaultdict(float)
    bulk = None
    for r, owner in zip(led, owners(target, led, refs)):
        if not owner:
            if r["source"] == "Debt sale":
                bulk = r
            continue
        x = loans.get(owner)
        if x is None:
            continue
        amt = float(r["amount_aud"])
        got_pv[owner] += amt * (1 + float(x["contract_rate"]) / 12) ** (-(mi(r["received_date"]) - mi(x["default_month"])))
    num, den = defaultdict(float), defaultdict(float)
    for x in rows:
        lno, B, r = x["loan_no"], float(x["balance_at_default"]), float(x["contract_rate"])
        disc = (1 + r / 12) ** (-(mi(bulk["received_date"]) - mi(x["default_month"])))
        sale = face.get(lno, 0.0) * price if sale_in else 0.0
        label = x["resolution"] == "Cured"
        cured = {"gold": label and lno not in gone, "register": label and lno not in stale, "label": label,
                 "no_withdrawn": label and lno not in withdrawn}[cures]
        loss = 0.0 if cured else B - got_pv[lno] - sale * disc
        q = ((int(x["default_month"][5:7]) - 7) % 12) // 3 + 1
        num[(x["product"], q)] += loss
        den[(x["product"], q)] += B
    return {kq: 100.0 * num[kq] / den[kq] for kq in den}


def dpd_clock(target, extract):
    """the dictionary's run clock for one loan book extract: (the close of business its loans are open at, the date its
    days_past_due are counted to), read from the extract's line in section A"""
    txt = " ".join(open(os.path.join(target, "ecl_extract_dictionary.txt")).read().split())
    m = re.search(re.escape(extract) + r" run [^:]+: every loan open at close of business on (\d{1,2} \w+ \d{4}); "
                  r"days_past_due counted to the [a-z ]+, (\d{1,2} \w+ \d{4})", txt)
    if m is None:
        return None
    f = lambda x: dt.datetime.strptime(x, "%d %B %Y").date()
    return f(m.group(1)), f(m.group(2))


def migration(target, june="dates", run="2026-07-01", dpd="close"):
    """the June extract's days past due are counted to the date its dictionary line gives: 'close' puts a loan in
    arrears back to close of business on 30 June (1.3) by the days between the two; 'extract' reads the column as it
    stands"""
    bk, hd, mult, conv = read_inputs(target)
    jb = pd.read_csv(os.path.join(target, "personal_loan_book_2026-06-30.csv"))
    if dpd == "close":
        close, counted = dpd_clock(target, "personal_loan_book_2026-06-30.csv")
        jb["days_past_due"] = np.maximum(jb["days_past_due"].to_numpy() - (counted - close).days, 0)
    sc = pd.read_parquet(os.path.join(target, "behaviour_score_runs_2026-03_2026-08.parquet"))
    sc["run_date"] = pd.to_datetime(sc["run_date"]).dt.strftime("%Y-%m-%d")
    pick = sc[sc["run_date"] == run]
    score = jb["loan_no"].map(dict(zip(pick["loan_no"], pick["score_v4"])))
    if june == "dates":
        hd2 = hd
    else:                                   # the register's labels as they stood at the September extract, or every cancellation out
        if june == "label":
            keep = hd["arrangement_status"] == "Active"
        else:
            keep = (hd["arrangement_start"] <= pd.Timestamp(JUNE)) & (hd["arrangement_end"] >= pd.Timestamp(JUNE)) & (hd["arrangement_status"] != "Cancelled")
        hd2 = hd[keep].assign(arrangement_start=pd.Timestamp("2000-01-01"), arrangement_end=pd.Timestamp("2099-12-31"),
                              cancelled_date=pd.NaT)
    jst = stages(jb, hd2, conv, JUNE, score)
    sst = stages(bk, hd, conv, REP, bk["current_score_v4"])
    now = dict(zip(bk["loan_no"], sst))
    out = {}
    for p in ("PL36", "PL60"):
        for s0 in (1, 2):
            ids = jb.loc[((jb["product"] == p) & (jst == s0)).to_numpy(), "loan_no"]
            v = [now.get(x) for x in ids]
            for s1 in (1, 2, 3):
                out[(p, s0, "stage %d" % s1)] = sum(1 for y in v if y == s1)
            out[(p, s0, "closed")] = sum(1 for y in v if y is None)
    return out


def main(target=T, golden=G):
    fails = []

    def chk(name, got, exp, tol=0):
        good = got is not None and abs(float(got) - float(exp)) <= tol
        print(("OK  " if good else "FAIL"), name, got, exp)
        if not good:
            fails.append(name)

    mil = lambda x: round(x / 1e6, 1)
    fig, st, bk = allowance(target)
    wb = openpyxl.load_workbook(os.path.join(golden, "ecl_allowance_2026Q3.xlsx"), data_only=True)
    ws = wb["Allowance"]
    chk("allowance (AUD million)", ws["B4"].value, mil(fig["total"]))
    chk("stage 1 allowance (AUD million)", ws["B5"].value, mil(fig["stage1"]))
    chk("stage 2 allowance (AUD million)", ws["B6"].value, mil(fig["stage2"]))
    chk("stage 1 loans", ws["B8"].value, int((st == 1).sum()))
    chk("stage 2 loans", ws["B9"].value, int((st == 2).sum()))
    chk("stage 3 loans", ws["B10"].value, int((st == 3).sum()))
    for how in ("drawdown", "km", "grade", "crude", "r4b"):
        x, _, _ = allowance(target, how)
        d = x["total"] / fig["total"] - 1
        far = mil(x["total"]) != mil(fig["total"])
        print(("OK  " if far else "FAIL"), "rival %-9s %14.0f %+7.2f%%, allowance %.1f against %.1f"
              % (how, x["total"], 100 * d, mil(x["total"]), mil(fig["total"])))
        if not far:
            fails.append("rival " + how)
    for how, label in (("own", "taken-over loans left out"), ("first", "months on book from the first repayment")):
        x, _, _ = allowance(target, how)
        same = all(mil(x[n]) == mil(fig[n]) for n in ("total", "stage1", "stage2"))
        print(("OK  " if same else "FAIL"), "%s %14.0f: all three figures in their bins" % (label, x["total"]))
        if not same:
            fails.append(label)
    L = lgd(target)
    quarter = {"Jul to Sep 2023": 1, "Oct to Dec 2023": 2, "Jan to Mar 2024": 3, "Apr to Jun 2024": 4}
    seen = 0
    for r in wb["Realised LGD"].iter_rows(min_row=4, values_only=True):
        if r[0] in ("PL36", "PL60"):
            seen += 1
            chk("realised loss %s %s" % (r[0], r[1]), r[4], round(L[(r[0], quarter[r[1]])], 1))
    chk("realised loss rows", seen, 8)
    for name, kw in (("cures at their word", dict(cures="label")), ("remittances as posted", dict(refs="posted")),
                     ("both", dict(cures="label", refs="posted")),
                     ("every withdrawn cure doubted", dict(cures="no_withdrawn", refs="posted")),
                     ("the sale left out", dict(sale_in=False))):
        x = lgd(target, **kw)
        apart = all(round(x[q], 1) != round(L[q], 1) for q in L)
        print(("OK  " if apart else "FAIL"), "ask A stop %-30s" % name, " ".join("%.1f" % x[q] for q in sorted(x)))
        if not apart:
            fails.append("ask A stop " + name)
    cj = dpd_clock(target, "personal_loan_book_2026-06-30.csv")
    cs = dpd_clock(target, "personal_loan_book_2026-09-30.csv")
    chk("June extract days past due counted days after its close", (cj[1] - cj[0]).days if cj else None, 1)
    chk("September extract days past due counted days after its close", (cs[1] - cs[0]).days if cs else None, 0)
    M = migration(target)
    seen = 0
    for r in wb["Stage migration"].iter_rows(min_row=4, max_row=7, values_only=True):
        for j, w in enumerate(("stage 1", "stage 2", "stage 3", "closed")):
            seen += 1
            chk("migration %s stage %s -> %s" % (r[0], r[1][-1], w), r[2 + j], M[(r[0], int(r[1][-1]), w)])
    chk("migration cells", seen, 16)
    for name, kw in (("days past due as extracted", dict(dpd="extract")),
                     ("labels, May run", dict(june="label", run="2026-06-01", dpd="extract")),
                     ("cancelled out, May run", dict(june="excl", run="2026-06-01", dpd="extract")),
                     ("dates, May run", dict(run="2026-06-01", dpd="extract")), ("cancelled out, June run", dict(june="excl"))):
        x = migration(target, **kw)
        moved = sum(x[c] != M[c] for c in M)
        print(("OK  " if moved >= 6 else "FAIL"), "ask B stop %-24s moves %d of 16" % (name, moved))
        if moved < 6:
            fails.append("ask B stop " + name)
    print("VERIFY", "FAIL" if fails else "PASS", len(fails))
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main(*sys.argv[1:3]) if len(sys.argv) > 2 else main())
