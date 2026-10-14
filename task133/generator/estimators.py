"""Every graded figure of task133, every rung of the main ladder and every stop on the two ask ladders, read from the
shipped files alone."""
import datetime as dt
import math
import os
import re

import numpy as np
import openpyxl
import pandas as pd
from pypdf import PdfReader

F = dict(book="personal_loan_book_2026-09-30.csv", june="personal_loan_book_2026-06-30.csv",
         record="loan_record_2026-09-30.parquet",
         hardship="hardship_register_2026-09-30.xlsx", scen="ifrs9_scenarios_2026Q3.xlsx",
         note="scorecard_v4_migration_note.pdf", scores="behaviour_score_runs_2026-03_2026-08.parquet",
         defaults="collections_defaults_2023-24.csv", receipts="collections_receipts_2023-07_2026-09.csv",
         sale="debt_sale_2026-03_schedule.xlsx", placements="wirraway_placements_2026-09-30.csv")
REPORT = dt.date(2026, 9, 30)
PRIOR = dt.date(2026, 6, 30)
SCEN_SHEET = "Q1 FY27 scenarios"
LGD = {"PL36": 0.70, "PL60": 0.76}
STAGE2_SCALAR = 2.4
QUARTERS = {1: "Jul to Sep 2023", 2: "Oct to Dec 2023", 3: "Jan to Mar 2024", 4: "Apr to Jun 2024"}
GO_LIVE = dt.date(2019, 10, 1)           # the loan system went live, taking over the earlier platform's loan book


def mi(d):
    """calendar month index of a date"""
    return d.year * 12 + d.month - 1


RM = mi(REPORT)


def _d(x):
    return dt.date.fromisoformat(str(x)[:10])


def months_due(first_rep, report=REPORT):
    if first_rep > report:
        return 0
    n = (report.year - first_rep.year) * 12 + (report.month - first_rep.month)
    if report.day >= first_rep.day:
        n += 1
    return n


# ------------------------------------------------------------------------------------------------- shared reads

def load_staging(target):
    """what staging needs: the September book, the hardship register and the v3 to v4 conversion"""
    bk = pd.read_csv(os.path.join(target, F["book"]))
    bk["first_rep"] = bk["first_repayment_date"].map(_d)
    hd = pd.read_excel(os.path.join(target, F["hardship"]), sheet_name=0)
    text = " ".join(p.extract_text() for p in PdfReader(os.path.join(target, F["note"])).pages)
    m = re.search(r"v4 = ([0-9.]+) x v3 - ([0-9.]+)", text)
    return bk, hd, (float(m.group(1)), -float(m.group(2)))


def load(target):
    bk = pd.read_csv(os.path.join(target, F["book"]))
    bk["first_rep"] = bk["first_repayment_date"].map(_d)
    hd = pd.read_excel(os.path.join(target, F["hardship"]), sheet_name=0)
    wb = openpyxl.load_workbook(os.path.join(target, F["scen"]), data_only=True)
    ws = wb[SCEN_SHEET]
    scen = []
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not r or r[0] is None:
            break
        scen.append((r[0], float(r[1]), float(r[2])))
    text = " ".join(p.extract_text() for p in PdfReader(os.path.join(target, F["note"])).pages)
    m = re.search(r"v4 = ([0-9.]+) x v3 - ([0-9.]+)", text)
    mapping = (float(m.group(1)), -float(m.group(2)))
    return bk, hd, scen, mapping


def record(target):
    """the loan record, each loan's months on book counted in calendar months from the month it was drawn (CRM-ECL-04
    3.2): x is the month on book of its first default (ev 'D'), of its closing without a default ('C'), or the month on
    book it has reached at the reporting date if it is still open and has not defaulted ('O'); entry is the last month on
    book before the record held the loan, 0 for the loan system's own loans and the month on book at 30 September 2019
    for a loan taken over from the earlier platform at the go-live"""
    rec = pd.read_parquet(os.path.join(target, F["record"]))
    draw = pd.to_datetime(rec["drawdown_date"])
    rec["draw_m"] = draw.dt.year * 12 + draw.dt.month - 1
    dm = rec["default_month"].map(lambda v: int(v[:4]) * 12 + int(v[5:7]) - 1 if isinstance(v, str) else -1)
    cd = pd.to_datetime(rec["closed_date"])
    cm = (cd.dt.year * 12 + cd.dt.month - 1).fillna(-1).astype(int)
    ev = np.where(dm >= 0, "D", np.where(cm >= 0, "C", "O"))
    em = np.where(dm >= 0, dm, np.where(cm >= 0, cm, RM))
    rec["ev"] = ev
    rec["x"] = em - rec["draw_m"]
    rec["entry"] = np.maximum(0, mi(GO_LIVE) - 1 - rec["draw_m"])
    rec["first_m"] = pd.to_datetime(rec["first_repayment_date"]).map(mi)
    return rec


def hazards(rec, entry="transfer", by=("product",), maxm=80, clock="calendar"):
    """each month on book's default rate and repayment rate: of the loans open and not in default at the start of the
    month on book that the record holds, the share that first default in it and the share that close without a default.
    entry 'transfer' counts a loan taken over at the go-live only from its first month on book after it; 'drawdown' counts
    every loan from its first month on book; 'own' leaves the taken-over loans out. clock 'first' counts months on book
    from the month of the first repayment instead of the drawdown month"""
    out = {}
    g0 = rec if entry != "own" else rec[rec["entry"] == 0]
    for key, g in g0.groupby(list(by)):
        shift = (g["first_m"] - g["draw_m"]).to_numpy() if clock == "first" else 0
        x = g["x"].to_numpy() - shift
        e = (g["entry"].to_numpy() - shift) if entry == "transfer" else np.full(len(g), -10 ** 6)
        t = g["ev"].to_numpy()
        d = np.zeros(maxm + 2)
        c = np.zeros(maxm + 2)
        n = np.zeros(maxm + 2)
        for m in range(1, maxm + 2):
            risk = (e < m) & (x >= m)
            k = int(risk.sum())
            n[m] = k
            if k:
                d[m] = ((x == m) & (t == "D") & risk).sum() / k
                c[m] = ((x == m) & (t == "C") & risk).sum() / k
        key = key if len(by) > 1 else (key[0] if isinstance(key, tuple) else key)
        out[key] = (d, c, n)
    return out


def landmark(rec, product, a, k):
    """the share of the product's loans the record shows open and not in default at the end of month on book a, among
    those drawn early enough to have been followed to a + k, that first defaulted in month on book a + k"""
    g = rec[rec["product"] == product]
    held = (g["entry"] <= a) & (g["x"] > a) & (RM - g["draw_m"] >= a + k)
    n = int(held.sum())
    return float(((g["x"] == a + k) & (g["ev"] == "D") & held).sum()) / n if n else 0.0


def in_force(hd, at, rule="dates"):
    """loans with a hardship arrangement in force on day `at`. 'dates' is CRM-ECL-04 2.2 (in force from its start date to
    its end date, both included, and not from the day it is cancelled); the others are the misreadings the ask ladders
    price"""
    at = pd.Timestamp(at)
    s, e, c = hd["arrangement_start"], hd["arrangement_end"], hd["cancelled_date"]
    status = hd["arrangement_status"]
    if rule == "dates":
        sel = (s <= at) & (e >= at) & (c.isna() | (c > at))
    elif rule == "excl_cancel":
        sel = (s <= at) & (e >= at) & (status != "Cancelled")
    elif rule == "ignore_cancel":
        sel = (s <= at) & (e >= at)
    elif rule == "label":
        sel = status == "Active"
    elif rule == "label_dated":
        sel = (status == "Active") & (s <= at)
    elif rule == "end_exclusive":
        sel = (s <= at) & (e > at) & (c.isna() | (c > at))
    elif rule == "cancel_inclusive":
        sel = (s <= at) & (e >= at) & (c.isna() | (c >= at))
    else:
        raise ValueError(rule)
    return set(hd.loc[sel, "loan_no"])


def stage(bk, hd, mapping, use_hardship=True, map_v3=True, at=REPORT, rule="dates", score=None):
    """CRM-ECL-04 section 2 at day `at`: stage 3 at 90 or more days past due; stage 2 at 30 or more, a hardship
    arrangement in force, or a behaviour score 60 or more points below the origination score on v4 (v3 converted)"""
    a, b = mapping
    o = bk["origination_score"].astype(float)
    if map_v3:
        o = [int(math.floor(a * v + b + 0.5)) if c == "v3" else int(v) for v, c in zip(o, bk["origination_scorecard"])]
    sc = bk["current_score_v4"] if score is None else score
    drop = pd.Series(o, index=bk.index) - pd.Series(sc, index=bk.index)
    st = pd.Series(1, index=bk.index)
    st[(bk["days_past_due"] >= 30) | (drop >= 60).fillna(False)] = 2
    if use_hardship:
        st[bk["loan_no"].isin(in_force(hd, at, rule))] = 2
    st[bk["days_past_due"] >= 90] = 3
    return st


# ------------------------------------------------------------------------------------------------- the main call

def ead_path(balance, rate, rem):
    """balance at the start of each month k = 1..rem, re-amortising the current balance over rem months"""
    i = rate / 12.0
    if rem <= 0:
        return []
    p = balance * i / (1 - (1 + i) ** -rem)
    out, b = [], balance
    for _ in range(rem):
        out.append(b)
        b = b * (1 + i) - p
    return out


def forward_pd(hz, a, H, competing=True, scale=1.0, on="pd"):
    """PD(k), k = 1..H, for a loan at month on book a: the default rate at month on book a + k, times the share of the
    product's loans that neither default nor close in each month on book a + 1 to a + k - 1. 'scale' is the stage 2
    scalar, on PD(k) (CRM-ECL-04 3.2) or, as a misreading, on the default rate inside the product"""
    d, c, _ = hz
    top = len(d) - 1
    out, surv = [], 1.0
    for k in range(1, H + 1):
        m = min(a + k, top)
        hd = d[m] * (scale if on == "hazard" else 1.0)
        out.append(surv * hd * (scale if on == "pd" else 1.0))
        surv *= 1.0 - hd - (c[m] if competing else 0.0)
    return out


def loan_ecl(row, st, hz, mult, competing=True, disc=True, flat_ead=False, pm_zero=False, scalar_on="pd", pd_fn=None,
             ead_runoff=False):
    """the loan's ECL under CRM-ECL-04 section 3, the PD from the record's rates (or pd_fn, a rung's own PD). ead_runoff
    also runs EAD(k) down by the share of the product's loans not yet repaid by month k, counting repayment a second
    time on top of the competing exit in the PD (the audit's double-count route)"""
    a = int(row["mob"])
    T = int(row["term_months"])
    rem = T - months_due(row["first_rep"])
    if rem <= 0:
        # past contractual maturity and not in stage 3: one month remaining (CRM-ECL-04 3.2)
        if pm_zero:
            return 0.0
        rem = 1
    H = rem if st == 2 else min(12, rem)
    r = float(row["contract_rate"])
    bal = float(row["current_balance"])
    eads = ead_path(bal, r, rem)
    scale = STAGE2_SCALAR if st == 2 else 1.0
    pds = pd_fn(row, a, H, scale) if pd_fn is not None else forward_pd(hz, a, H, competing, scale, scalar_on)
    tot = 0.0
    stay = 1.0
    for k in range(1, H + 1):
        ead = bal if flat_ead else eads[k - 1]
        if ead_runoff:
            ead *= stay
            stay *= 1.0 - hz[1][min(a + k, len(hz[1]) - 1)]
        df = (1 + r / 12.0) ** (-k) if disc else 1.0
        tot += pds[k - 1] * mult * LGD[row["product"]] * ead * df
    return tot


def crude_curves(rec, closed_only=True):
    """the share of every loan drawn that had defaulted (D) or closed without a default (Cl) by the end of each month on
    book, over the loans closed by the reporting date (closed_only) or over every loan in the record"""
    out = {}
    for p, g in rec.groupby("product"):
        if closed_only:
            g = g[g["ev"] != "O"]
        n = len(g)
        grid = np.arange(0, 82)
        dm = np.sort(g.loc[g["ev"] == "D", "x"].to_numpy())
        cm = np.sort(g.loc[g["ev"] == "C", "x"].to_numpy())
        out[p] = (np.searchsorted(dm, grid, side="right") / n, np.searchsorted(cm, grid, side="right") / n)
    return out


def crude_pd_fn(cv, variant):
    def f(row, a, H, scale):
        C, Cl = cv[row["product"]]
        top = len(C) - 1
        out = []
        for k in range(1, H + 1):
            m = min(a + k, top)
            inc = C[m] - C[m - 1]
            if variant == "r4":
                v = inc / max(1e-12, 1.0 - C[min(a, top)])
            elif variant == "S":
                v = inc / max(1e-12, 1.0 - C[min(a, top)] - Cl[min(a, top)])
            elif variant == "shift":
                v = inc
            elif variant == "uncond":
                v = C[min(k, top)] - C[min(k, top) - 1]
            else:
                raise ValueError(variant)
            out.append(v * scale)
        return out
    return f


def landmark_pd_fn(rec):
    cache = {}

    def f(row, a, H, scale):
        out = []
        for k in range(1, H + 1):
            key = (row["product"], a, k)
            if key not in cache:
                cache[key] = landmark(rec, row["product"], a, k)
            out.append(cache[key] * scale)
        return out
    return f


def solve(target, **kw):
    """the allowance and its parts. kw: entry ('transfer', 'drawdown', 'own'), competing, by, clock, scalar_on, disc,
    flat_ead, pm_zero, pd ('hazard', 'landmark', 'crude_r4', 'crude_S', 'crude_shift', 'crude_uncond', each with
    closed_only), the staging knobs (use_hardship, map_v3, rule, stage_fn), mult, rec (a record already read)"""
    bk, hd, scen, mapping = load(target)
    rec = kw.get("rec")
    if rec is None:
        rec = record(target)
    st = stage(bk, hd, mapping, kw.get("use_hardship", True), kw.get("map_v3", True), rule=kw.get("rule", "dates"))
    if kw.get("stage_fn") is not None:
        st = kw["stage_fn"](bk, hd, mapping)
    mult = kw.get("mult", sum(w * m for _, w, m in scen))
    by = tuple(kw.get("by", ("product",)))
    hz = hazards(rec, kw.get("entry", "transfer"), by, clock=kw.get("clock", "calendar"))
    how = kw.get("pd", "hazard")
    pd_fn = None
    if how == "landmark":
        pd_fn = landmark_pd_fn(rec)
    elif how.startswith("crude_"):
        pd_fn = crude_pd_fn(crude_curves(rec, kw.get("closed_only", True)), how[6:])
    dm = pd.to_datetime(bk["first_repayment_date"] if kw.get("clock", "calendar") == "first" else bk["drawdown_date"])
    bk = bk.assign(mob=RM - (dm.dt.year * 12 + dm.dt.month - 1))
    ecl = []
    for idx, row in bk.iterrows():
        s = int(st[idx])
        if s == 3:
            ecl.append(0.0)
            continue
        key = tuple(row[b] for b in by) if len(by) > 1 else row[by[0]]
        ecl.append(loan_ecl(row, s, hz.get(key), mult, kw.get("competing", True), kw.get("disc", True),
                            kw.get("flat_ead", False), kw.get("pm_zero", False), kw.get("scalar_on", "pd"), pd_fn,
                            kw.get("ead_runoff", False)))
    bk = bk.assign(stage=st.values, ecl=ecl, due=[months_due(f) for f in bk["first_rep"]])
    inb = bk[bk["stage"] < 3]
    out = dict(total=float(inb["ecl"].sum()), mult=mult,
               loans=inb[["loan_no", "product", "stage", "current_balance", "ecl"]].copy(), hz=hz)
    for s_ in (1, 2):
        g = inb[inb["stage"] == s_]
        out["stage%d" % s_] = float(g["ecl"].sum())
        out["n%d" % s_] = int(len(g))
        out["gca%d" % s_] = float(g["current_balance"].sum())
    out["n3"] = int((bk["stage"] == 3).sum())
    pm = (bk["due"] >= bk["term_months"]) & (bk["stage"] < 3)
    out["n_pm"] = int(pm.sum())
    out["ecl_pm"] = float(bk.loc[pm, "ecl"].sum())
    out["stages"] = dict(zip(bk["loan_no"], bk["stage"]))
    return out


# ------------------------------------------------------------------------------------------------- ask A

_CACHE = {}


def _cached(key, make):
    """files read once per target, handed out as copies"""
    if key not in _CACHE:
        _CACHE[key] = make()
    v = _CACHE[key]
    return tuple(x.copy() if hasattr(x, "copy") else x for x in v) if isinstance(v, tuple) else v.copy()


def collections(target):
    return _cached(("collections", target), lambda: _collections(target))


def _collections(target):
    fy = pd.read_csv(os.path.join(target, F["defaults"]), dtype={"loan_no": str})
    led = pd.read_csv(os.path.join(target, F["receipts"]), dtype={"loan_no": str, "receipt_ref": str})
    wb = openpyxl.load_workbook(os.path.join(target, F["sale"]), data_only=True)
    ws = wb.worksheets[0]
    head, rows, cols = {}, [], None
    for r in ws.iter_rows(values_only=True):
        if cols is None:
            if r[0] == "loan_no":
                cols = list(r)
            elif r[0]:
                head[r[0]] = r[1]
            continue
        if r[0]:
            rows.append(dict(zip(cols, r)))
    sale = pd.DataFrame(rows)
    price = float(re.match(r"([0-9.]+) cents", str(head["Price"])).group(1)) / 100.0
    return fy, led, sale, price, head


def _month(s):
    s = str(s)
    return int(s[:4]) * 12 + int(s[5:7]) - 1


def placements(target):
    return _cached(("placements", target), lambda: _placements(target))


def _placements(target):
    pl = pd.read_csv(os.path.join(target, F["placements"]), dtype={"Client Ref": str})
    pl["placed"] = pd.to_datetime(pl["Placed"], format="%d/%m/%Y")
    pl["closed"] = pd.to_datetime(pl["Closed"], format="%d/%m/%Y")
    return pl


def attribute(target, led, refs="asof"):
    """the loan each receipt counts against. An agency remittance quotes Wirraway's reference in its narrative: 'asof'
    gives it to the placement that carried that reference on the day it was received; 'posted' takes the ledger's loan
    number as it stands; 'latest' gives it to the last placement to carry the reference; 'drop' leaves out every
    remittance quoting a reference more than one placement has carried"""
    led = led.copy()
    if refs == "posted":
        return led
    pl = placements(target)
    ag = led["source"] == "Agency remittance"
    led["_ref"] = led["narrative"].str.extract(r"ref (\d+)$")[0].astype(float)
    if refs == "drop":
        many = set(pl["Agency Ref"].value_counts().loc[lambda v: v > 1].index)
        return led[~(ag & led["_ref"].isin(many))].drop(columns="_ref")
    if refs == "latest":
        last = pl.sort_values("placed").groupby("Agency Ref")["Client Ref"].last()
        led.loc[ag, "loan_no"] = led.loc[ag, "_ref"].map(last).values
        return led.drop(columns="_ref")
    if refs != "asof":
        raise ValueError(refs)
    a = led.loc[ag, ["_ref", "received_date"]].reset_index()
    a["day"] = pd.to_datetime(a["received_date"])
    m = a.merge(pl[["Agency Ref", "Client Ref", "placed", "closed"]], left_on="_ref", right_on="Agency Ref", how="left")
    m = m[(m["placed"] <= m["day"]) & (m["closed"].isna() | (m["day"] <= m["closed"]))]
    assert m["index"].is_unique and len(m) == len(a), "each remittance falls inside one placement of its reference"
    led.loc[m["index"].values, "loan_no"] = m["Client Ref"].values
    return led.drop(columns="_ref")


def withdrawn_cures(target, fy, led):
    """the 2023-24 defaults Collections closed as cured on a collections hardship arrangement that was withdrawn before it
    started: the register holds an arrangement received after the default month end, starting on the case's resolution
    date and cancelled before that start. Returns those loans, and the ones among them that never made good their
    arrears: the loan record closes them as written off or settled for less, where a loan that caught up stayed open or
    was later repaid in full"""
    hd = _cached(("hardship", target), lambda: pd.read_excel(os.path.join(target, F["hardship"]), sheet_name=0))
    cur = fy[fy["resolution"] == "Cured"][["loan_no", "default_month", "resolution_date"]].copy()
    cur["dme"] = pd.to_datetime(cur["default_month"] + "-01") + pd.offsets.MonthEnd(0)
    cur["res"] = pd.to_datetime(cur["resolution_date"])
    x = hd.merge(cur, on="loan_no")
    w = x[(x["received_date"] > x["dme"]) & (x["arrangement_start"] == x["res"]) & (x["arrangement_status"] == "Cancelled")
          & (x["cancelled_date"] < x["arrangement_start"])]
    withdrawn = set(w["loan_no"])
    rec = _cached(("closures", target), lambda: pd.read_parquet(os.path.join(target, F["record"]), columns=["loan_no", "closure_reason"]))
    gone = set(rec.loc[rec["closure_reason"].isin(["Written off", "Settled"]), "loan_no"])
    return withdrawn, withdrawn & gone


def realised_lgd(target, variant="gold", cures="gold", refs="asof"):
    """what the 2023-24 defaults have lost, as a share of their balance at default, by product and quarter of the
    default month, in per cent. 'gold' is CRM-ECL-04 3.3: every amount received after default to the reporting date, the
    loan's share of the March 2026 sale proceeds (its face value at the cut-off times the price) included, counted at the
    end of the month it was received and discounted to the default month end monthly at the contract rate; a loan that
    returned to performing has lost nothing. cures: 'gold' takes a loan as returned to performing when Collections closed
    it as cured, unless the arrangement it was closed on never started and the loan was written off; 'label' takes every
    case closed as cured at its word; 'no_withdrawn' takes none closed on a withdrawn arrangement (the over-cleaner).
    refs: how agency remittances are given to loans (attribute). The other variants are the stops on the ask's ladder"""
    fy, led, sale, price, head = collections(target)
    led = attribute(target, led, refs)
    fy = fy.copy()
    fy["B"] = fy["balance_at_default"].astype(float)
    fy["r"] = fy["contract_rate"].astype(float)
    fy["M"] = fy["default_month"].map(_month)
    fy["q"] = [((int(m[5:7]) - 7) % 12) // 3 + 1 for m in fy["default_month"]]
    rec = led[led["loan_no"].notna()].merge(fy[["loan_no", "M", "r", "default_month"]], on="loan_no", how="inner")
    rec["amt"] = rec["amount_aud"].astype(float)
    rec["n"] = rec["received_date"].map(_month) - rec["M"]
    me = pd.to_datetime(rec["default_month"] + "-01") + pd.offsets.MonthEnd(0)
    rec["days"] = (pd.to_datetime(rec["received_date"]) - me).dt.days
    rec["pv"] = rec["amt"] * (1 + rec["r"] / 12) ** (-rec["n"])
    rec["pv_days"] = rec["amt"] * (1 + rec["r"] / 12) ** (-rec["days"] / (365.25 / 12))
    rec["pv_start"] = rec["amt"] * (1 + rec["r"] / 12) ** (-(rec["n"] - 1))
    rec["pv_annual"] = rec["amt"] * (1 + rec["r"]) ** (-rec["n"] / 12.0)
    rec["pv_kept"] = np.where(rec["amt"] < 0, 0.0, rec["pv"])
    by = rec.groupby("loan_no")
    for c in ("amt", "pv", "pv_days", "pv_start", "pv_annual", "pv_kept"):
        fy[c] = fy["loan_no"].map(by[c].sum()).fillna(0.0)
    bulk = led[led["loan_no"].isna() & (led["source"] == "Debt sale")]
    sale_m = _month(bulk["received_date"].iloc[0])
    sale_day = pd.Timestamp(bulk["received_date"].iloc[0])
    consideration = float(bulk["amount_aud"].astype(float).sum())
    face = dict(zip(sale["loan_no"], sale["face_value_aud"].astype(float)))
    fy["sale"] = fy["loan_no"].map(face).fillna(0.0) * price
    fy["sale_pv"] = fy["sale"] * (1 + fy["r"] / 12) ** (-(sale_m - fy["M"]))
    dme = pd.to_datetime(fy["default_month"] + "-01") + pd.offsets.MonthEnd(0)
    fy["sale_pv_days"] = fy["sale"] * (1 + fy["r"] / 12) ** (-((sale_day - dme).dt.days) / (365.25 / 12))
    fy["sale_pv_start"] = fy["sale"] * (1 + fy["r"] / 12) ** (-(sale_m - fy["M"] - 1))
    fy["sale_pv_annual"] = fy["sale"] * (1 + fy["r"]) ** (-(sale_m - fy["M"]) / 12.0)
    withdrawn, written_off = withdrawn_cures(target, fy, led)
    label = fy["resolution"] == "Cured"
    if cures == "gold":
        cured = label & ~fy["loan_no"].isin(written_off)
    elif cures == "label":
        cured = label
    elif cures == "no_withdrawn":
        cured = label & ~fy["loan_no"].isin(withdrawn)
    else:
        raise ValueError(cures)
    wo = fy["resolution"] == "Written off"
    pop = None
    if variant == "gold":
        loss = np.where(cured, 0.0, fy["B"] - fy["pv"] - fy["sale_pv"])
    elif variant == "s1":            # cases closed as written off only, undiscounted, the sale left out
        loss, pop = fy["B"] - fy["amt"], wo
    elif variant == "s2":            # cures in at full recovery, undiscounted, the sale left out
        loss = np.where(cured, 0.0, fy["B"] - fy["amt"])
    elif variant == "s3":            # discounted, the sale left out
        loss = np.where(cured, 0.0, fy["B"] - fy["pv"])
    elif variant == "s4":            # the sale receipt spread over the 2023-24 written-off loans by balance at default
        share = np.where(wo, fy["B"] / fy.loc[wo, "B"].sum(), 0.0)
        loss = np.where(cured, 0.0, fy["B"] - fy["pv"] - consideration * share * (1 + fy["r"] / 12) ** (-(sale_m - fy["M"])))
    elif variant == "cures_receipts":
        loss = fy["B"] - fy["pv"] - fy["sale_pv"]
    elif variant == "excl_cures":
        loss, pop = fy["B"] - fy["pv"] - fy["sale_pv"], ~cured
    elif variant == "sale_undisc":
        loss = np.where(cured, 0.0, fy["B"] - fy["pv"] - fy["sale"])
    elif variant == "all_undisc":
        loss = np.where(cured, 0.0, fy["B"] - fy["amt"] - fy["sale"])
    elif variant == "floor":
        loss = np.maximum(np.where(cured, 0.0, fy["B"] - fy["pv"] - fy["sale_pv"]), 0.0)
    elif variant == "days":
        loss = np.where(cured, 0.0, fy["B"] - fy["pv_days"] - fy["sale_pv_days"])
    elif variant == "month_start":
        loss = np.where(cured, 0.0, fy["B"] - fy["pv_start"] - fy["sale_pv_start"])
    elif variant == "cure_disc":      # a cure read as recovering its balance at the cure, discounted to the default
        nres = pd.to_datetime(fy["resolution_date"]).dt.year * 12 + pd.to_datetime(fy["resolution_date"]).dt.month - 1
        loss = np.where(cured, fy["B"] * (1 - (1 + fy["r"] / 12) ** (-(nres - fy["M"]))), fy["B"] - fy["pv"] - fy["sale_pv"])
    elif variant == "annual":         # the annual contract rate compounded per whole month, (1 + r) ** (-n / 12)
        loss = np.where(cured, 0.0, fy["B"] - fy["pv_annual"] - fy["sale_pv_annual"])
    elif variant == "lazy":           # every receipt summed as it stands, cures taken from their receipts, no sale
        loss = fy["B"] - fy["amt"]
    elif variant == "no_reversals":   # the negative ledger rows dropped as errors, the dishonoured attempts kept
        loss = np.where(cured, 0.0, fy["B"] - fy["pv_kept"] - fy["sale_pv"])
    elif variant == "face_alloc":     # the sale receipt shared over the sold loans in proportion to face value
        alloc = fy["loan_no"].map(face).fillna(0.0) / sum(face.values()) * consideration
        loss = np.where(cured, 0.0, fy["B"] - fy["pv"] - alloc * (1 + fy["r"] / 12) ** (-(sale_m - fy["M"])))
    else:
        raise ValueError(variant)
    d = fy.assign(loss=loss)
    if pop is not None:
        d = d[pop]
    g = d.groupby(["product", "q"])
    val = 100.0 * g["loss"].sum() / g["B"].sum()
    return {(p, int(q)): float(v) for (p, q), v in val.items()}


def lgd_path(target, cures="gold", sale="share", disc="monthly", refs="asof"):
    """the ask A figures along one path, each knob a way of handling one device or hazard: cures 'gold', 'label' (every
    case closed as cured at its word), 'no_withdrawn' (none closed on a withdrawn arrangement) or 'receipts' (every loan on
    its receipts); sale 'share' (each sold loan's face value times the price), 'none', 'spread' (the receipt over the
    2023-24 written-off cases by balance) or 'flat' (shares not discounted); disc 'monthly' (3.3 and 3.5), 'none',
    'annual' ((1 + r) ** (-n / 12)), 'days' or 'month_start'; refs as in attribute"""
    fy, led, sale_sch, price, head = collections(target)
    led = attribute(target, led, refs)
    fy = fy.copy()
    fy["B"] = fy["balance_at_default"].astype(float)
    fy["r"] = fy["contract_rate"].astype(float)
    fy["M"] = fy["default_month"].map(_month)
    fy["q"] = [((int(m[5:7]) - 7) % 12) // 3 + 1 for m in fy["default_month"]]
    rec = led[led["loan_no"].notna()].merge(fy[["loan_no", "M", "r", "default_month"]], on="loan_no", how="inner")
    rec["amt"] = rec["amount_aud"].astype(float)
    rec["n"] = rec["received_date"].map(_month) - rec["M"]
    me = pd.to_datetime(rec["default_month"] + "-01") + pd.offsets.MonthEnd(0)
    days = (pd.to_datetime(rec["received_date"]) - me).dt.days

    def factor(n, r, dd):
        if disc == "monthly":
            return (1 + r / 12) ** (-n)
        if disc == "none":
            return 1.0 + 0 * n
        if disc == "annual":
            return (1 + r) ** (-n / 12.0)
        if disc == "days":
            return (1 + r / 12) ** (-dd / (365.25 / 12))
        if disc == "month_start":
            return (1 + r / 12) ** (-(n - 1))
        raise ValueError(disc)
    rec["pv"] = rec["amt"] * factor(rec["n"], rec["r"], days)
    fy["pv"] = fy["loan_no"].map(rec.groupby("loan_no")["pv"].sum()).fillna(0.0)
    bulk = led[led["loan_no"].isna() & (led["source"] == "Debt sale")]
    sale_m = _month(bulk["received_date"].iloc[0])
    sale_day = pd.Timestamp(bulk["received_date"].iloc[0])
    dme = pd.to_datetime(fy["default_month"] + "-01") + pd.offsets.MonthEnd(0)
    f_sale = factor(sale_m - fy["M"], fy["r"], (sale_day - dme).dt.days)
    face = dict(zip(sale_sch["loan_no"], sale_sch["face_value_aud"].astype(float)))
    wo = fy["resolution"] == "Written off"
    if sale == "share":
        fy["sale_pv"] = fy["loan_no"].map(face).fillna(0.0) * price * f_sale
    elif sale == "flat":
        fy["sale_pv"] = fy["loan_no"].map(face).fillna(0.0) * price
    elif sale == "none":
        fy["sale_pv"] = 0.0
    elif sale == "spread":
        cons = float(bulk["amount_aud"].astype(float).sum())
        fy["sale_pv"] = np.where(wo, fy["B"] / fy.loc[wo, "B"].sum(), 0.0) * cons * f_sale
    else:
        raise ValueError(sale)
    withdrawn, written_off = withdrawn_cures(target, fy, led)
    label = fy["resolution"] == "Cured"
    cured = {"gold": label & ~fy["loan_no"].isin(written_off), "label": label,
             "no_withdrawn": label & ~fy["loan_no"].isin(withdrawn), "receipts": label & False}[cures]
    loss = np.where(cured, 0.0, fy["B"] - fy["pv"] - fy["sale_pv"])
    g = fy.assign(loss=loss).groupby(["product", "q"])
    val = 100.0 * g["loss"].sum() / g["B"].sum()
    return {(p, int(q)): float(v) for (p, q), v in val.items()}


# ------------------------------------------------------------------------------------------------- ask B

def score_runs(target):
    sc = pd.read_parquet(os.path.join(target, F["scores"]))
    sc["run_date"] = pd.to_datetime(sc["run_date"]).dt.date
    return sc


def migration(target, june_rule="dates", run="june", sept_rule="dates", june_dpd="close"):
    """the loans in stage 1 and stage 2 at 30 June 2026, by product and June stage: how many are in stage 1, 2 and 3
    at 30 September and how many have closed (not in the September extract). 'june' takes the scores of the month-end
    run on the 30 June position (made on 1 July); 'may' the last run made on or before 30 June; 'sept' the September
    extract's scores. june_dpd 'close' takes days past due at close of business on 30 June (CRM-ECL-04 1.3): the June
    extract, run at 07:30 on 1 July before CHG-26-0418, counts them to the run, a day more; 'extract' takes the column
    as it stands"""
    bk, hd, mapping = load_staging(target)
    jb = pd.read_csv(os.path.join(target, F["june"]))
    if june_dpd == "close":
        jb["days_past_due"] = (jb["days_past_due"] - 1).clip(lower=0)
    elif june_dpd != "extract":
        raise ValueError(june_dpd)
    sc = score_runs(target)
    if run == "june":
        pick = sc[sc["run_date"] == dt.date(2026, 7, 1)]
    elif run == "may":
        pick = sc[sc["run_date"] == max(d for d in sc["run_date"].unique() if d <= PRIOR)]
    elif run == "july":
        pick = sc[sc["run_date"] == dt.date(2026, 8, 1)]
    elif run == "sept":
        pick = bk[["loan_no", "current_score_v4"]].rename(columns={"current_score_v4": "score_v4"})
    else:
        raise ValueError(run)
    score = jb["loan_no"].map(dict(zip(pick["loan_no"], pick["score_v4"])))
    jst = stage(jb, hd, mapping, at=PRIOR, rule=june_rule, score=score)
    sst = stage(bk, hd, mapping, rule=sept_rule)
    now = dict(zip(bk["loan_no"], sst))
    cells = {}
    for p in ("PL36", "PL60"):
        for s0 in (1, 2):
            ids = jb.loc[(jb["product"] == p) & (jst == s0), "loan_no"]
            ss = ids.map(now)
            for s1 in (1, 2, 3):
                cells[(p, s0, "stage %d" % s1)] = int((ss == s1).sum())
            cells[(p, s0, "closed")] = int(ss.isna().sum())
    return cells
