"""Assertions on task133's shipped pack: the input gates; the three rules (R1 to R3c); the allowance, every rung and rival
of the main ladder and every reading the fork table admits (SF1 to SF18); the loan record and its take-over structure;
the realism of the extracts (repayment calendar, balances, hardship register, pricing, scores); the June extract and the
month-end scoring runs; the Collections files, the March 2026 sale and the agency's placements; both ask ladders (SF19 to
SF28); the separation of the ask devices from the main call's rows; the pins; the ask-layer audits (lazy sweep, necessity
matrix, composed mishandlings, pair simulation, over-cleaner, oracle sweep, hygiene battery); and the deliverables. It
writes generator/figures.json for figures_check.py.

    python3 checks.py [target golden]
"""
import datetime as dt
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
import zipfile
from collections import Counter

import numpy as np
import pandas as pd
import pyarrow.parquet as pq

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import estimators as E
import officex as OX
import verify as V
import world as W

TASK = os.path.dirname(HERE)
T = os.path.join(TASK, "target")
G = os.path.join(TASK, "golden")
if len(sys.argv) > 2:
    T, G = os.path.abspath(sys.argv[1]), os.path.abspath(sys.argv[2])
R = W.REPORT
JUNE = W.PRIOR_QE
RERUN = dt.date(2026, 7, 2)                # the June month end's partition, run again
n_ok = 0
LOG = []


def ok(cond, msg):
    global n_ok
    if not cond:
        raise AssertionError(msg)
    n_ok += 1


def say(*a):
    line = " ".join(str(x) for x in a)
    LOG.append(line)
    print(line)


def edge(v, unit):
    """distance from the nearest rounding edge, in bins (0.5 is mid-bin)"""
    return abs((v / unit) % 1 - 0.5)


M = lambda x: round(x / 1e6, 1)              # the allowance and its stage amounts as reported: AUD million, one decimal
FIG = ("total", "stage1", "stage2")
HOL = set()
for _y in range(2014, 2033):
    HOL |= W.nsw_holidays(_y)


def is_bd(d):
    return d.weekday() < 5 and d not in HOL


def due_dates(first, n):
    out = []
    for j in range(n):
        y, m = first.year + (first.month - 1 + j) // 12, (first.month - 1 + j) % 12 + 1
        d = dt.date(y, m, first.day)
        while not is_bd(d):
            d += dt.timedelta(days=1)
        out.append(d)
    return out


def docx_text(path):
    with zipfile.ZipFile(path) as z:
        return " ".join(re.sub(r"<[^>]+>", "", p) for p in re.findall(r"<w:p [^>]*>.*?</w:p>", z.read("word/document.xml").decode(), re.S))


def meta_of(path):
    return {k.decode(): v.decode() for k, v in pq.read_schema(path).metadata.items() if not k.startswith(b"pandas")}


def POLICY():
    return os.path.join(T, "CRM-ECL-04_expected_credit_loss_personal_loans.docx")


ASK_A_FILES = ["defaults", "receipts", "sale", "placements"]
ASK_B_FILES = ["june", "scores"]
DISTRACTORS = ["write_offs_2025-07_2026-09.csv", "arrears_roll_rates_fy26.xlsx"]


# ---------------------------------------------------------------------------------------------------------- gates

def gates():
    files = sorted(os.listdir(T))
    fmts = {f.rsplit(".", 1)[1] for f in files}
    ok(len(files) >= 10, "10 or more input files (%d)" % len(files))
    ok(len(fmts) >= 3, "3 or more formats (%s)" % sorted(fmts))
    rows = {f: len(pd.read_csv(os.path.join(T, f), dtype=str)) for f in files if f.endswith(".csv")}
    rows.update({f: pq.read_metadata(os.path.join(T, f)).num_rows for f in files if f.endswith(".parquet")})
    ok(max(rows.values()) >= 25000, "a file of 25,000 or more rows (%s)" % {k: v for k, v in rows.items() if v >= 25000})
    meta = json.load(open(os.path.join(TASK, "metadata.json")))
    ok(len(meta["distractors"]) >= 2 and all(d in files for d in meta["distractors"]), "two or more distractors on disk")
    ok(set(meta["distractors"]) == set(DISTRACTORS), "the distractors are the roll-rate workbook and the aggregate write-off file")
    ok(sorted(x["path"] for x in meta["input_files"]) == files, "metadata lists every shipped file")
    ok(meta["deliverables"] == ["ecl_allowance_2026Q3.xlsx", "arc_ecl_q3_2026.pptx"]
       and all(os.path.exists(os.path.join(G, d)) for d in meta["deliverables"]), "two deliverables, both built")
    for f in files:
        ok(not any(w in f.lower() for w in ("answer", "golden", "trap", "clean", "_v2", "test_", "distract", "legacy",
                                            "truncat", "reissu")), "file name %s carries no author noun" % f)
    return len(files), len(fmts), rows


# ---------------------------------------------------------------------------------------------------------- the loan record

def loan_record():
    """the record the PD is read from (CRM-ECL-04 3.2): one row per loan the loan system has held since it went live on
    1 October 2019, the loans it took over from the earlier platform included, each with its drawdown, first default
    month and closing; no column says which loans came across or how long the record has held each"""
    rec = pd.read_parquet(os.path.join(T, E.F["record"]))
    ok(list(rec.columns) == ["loan_no", "product", "risk_grade", "drawdown_date", "first_repayment_date", "term_months",
                             "contract_rate", "original_amount", "default_month", "status", "closed_date", "closure_reason"],
       "R1: the record ships each loan's dates and outcome at loan grain, and no month on book, rate, flag or PD column")
    ok(rec["loan_no"].is_unique and rec["loan_no"].is_monotonic_increasing, "the record is one row per loan, in loan number order")
    meta = meta_of(os.path.join(T, E.F["record"]))
    ok("since it went live on 1 October 2019, when it took over the personal loan book from the earlier platform" in meta["population"]
       and all(c in meta for c in rec.columns), "the record's population names the go-live take-over; every column is described")
    draw = pd.to_datetime(rec["drawdown_date"])
    old = draw < pd.Timestamp(W.GO_LIVE)
    closed = pd.to_datetime(rec["closed_date"])
    ok((closed[old].dropna() >= pd.Timestamp(W.GO_LIVE)).all() and (rec.loc[old, "default_month"].dropna() >= "2019-10").all(),
       "every loan drawn before the go-live closed and defaulted, if at all, on the loan system: the earlier platform's "
       "closed loans and defaults before October 2019 never came across (%d loans drawn before it)" % int(old.sum()))
    ok(rec.loc[old, "status"].eq("Closed").all(), "no loan taken over at the go-live is still open")
    yrs = draw.dt.year.value_counts().sort_index()
    ok(yrs[2014] < 0.05 * yrs[2018] and yrs.loc[2015:2018].is_monotonic_increasing and yrs.loc[2020:2025].is_monotonic_increasing,
       "the record's early vintages are thin, only the loans still open in 2019 (%s)" % yrs.loc[2014:2019].to_dict())
    pl60_old = rec.loc[old & (draw.dt.year <= 2016), "product"].value_counts()
    ok(pl60_old.get("PL60", 0) > 3 * pl60_old.get("PL36", 0), "of the 2014 to 2016 loans the record holds mostly PL60s, the "
       "long loans still running at the go-live (%s)" % pl60_old.to_dict())
    num = rec["loan_no"].str[3:9].astype(int)
    ok(num[old].max() < num[~old].min(), "the earlier platform's loan numbers sit below the loan system's own")
    bk = pd.read_csv(os.path.join(T, E.F["book"]), dtype={"loan_no": str})
    ok(set(rec.loc[rec["status"] == "Open", "loan_no"]) == set(bk["loan_no"]), "SF2: the record's open loans are exactly the "
       "September extract's (%d)" % len(bk))
    common = rec.merge(bk, on="loan_no", suffixes=("", "_b"))
    for c in ("product", "risk_grade", "drawdown_date", "first_repayment_date", "term_months", "original_amount"):
        ok((common[c].astype(str) == common[c + "_b"].astype(str)).all(), "the record and the September extract agree on %s" % c)
    cl = rec[rec["status"] == "Closed"]
    ok(cl["closed_date"].notna().all() and cl["closure_reason"].notna().all() and rec.loc[rec["status"] == "Open", "closed_date"].isna().all(),
       "closed loans carry a closing date and reason, open ones neither")
    wo = rec[rec["closure_reason"] == "Written off"]
    wc = pd.to_datetime(wo["closed_date"])
    ok(wo["default_month"].notna().all() and (wc == wc + pd.offsets.MonthEnd(0)).all(),
       "every written-off loan defaulted first and closed at a month end (the write-off batch)")
    gap = (wc.dt.to_period("M") - pd.to_datetime(wo["default_month"]).dt.to_period("M")).map(lambda x: x.n)
    ok(gap.min() >= 2 and gap.median() == 3, "write-offs follow the default by about three month ends, 90 to 180 days past due")
    st = rec[rec["closure_reason"] == "Settled"]
    ok(len(st) > 100 and st["default_month"].notna().mean() > 0.6, "settlements for less mostly follow a recorded default; the "
       "rest settled inside the month they reached 90 days (%d of %d)" % (int(st["default_month"].isna().sum()), len(st)))
    rate = rec.assign(D=rec["default_month"].notna()).groupby("risk_grade")["D"].mean()
    ok(rate.is_monotonic_increasing and rate["A"] < 0.04 and rate["E"] < 0.2, "default shares rise A to E at a sustainable level (%s)"
       % rate.round(3).to_dict())
    rec2 = E.record(T)
    took = rec2["entry"] > 0
    ok(int(took.sum()) > 10000 and int((took & (rec2["ev"] == "D")).sum()) > 500,
       "the take-over brought %d loans the record holds only from their month on book at 30 September 2019, %d of them "
       "later in default" % (int(took.sum()), int((took & (rec2["ev"] == "D")).sum())))
    return dict(rows=len(rec), taken_over=int(took.sum()), drawn_before=int(old.sum()), open=int((rec["status"] == "Open").sum()),
                defaults=int(rec["default_month"].notna().sum()))


# ---------------------------------------------------------------------------------------------------------- the main call

# the ladder and every rival measured on the allowance. 'rung': a defensible analysis landing off the answer's bin;
# 'c4': a cell a quoted sentence rules out, landing off the bin; 'same': a reading or method that lands every figure in the
# answer's bin
MAIN = [
    ("R0 default path: every loan counted from its drawdown", dict(entry="drawdown"), "rung", None),
    ("R1 the reflex: risk grade added to the rates, from drawdown", dict(entry="drawdown", by=("product", "risk_grade")), "rung", None),
    ("R2 the textbook survival curve: repayment censored, from drawdown", dict(entry="drawdown", competing=False), "rung", None),
    ("R2b repayment censored, held from the take-over", dict(competing=False), "rung", "SF4"),
    ("audit route: crude cumulative default share of the resolved loans, over 1 - C(a)", dict(pd="crude_r4"), "rung", None),
    ("audit route: crude share over the share neither defaulted nor closed", dict(pd="crude_S"), "rung", None),
    ("audit route: crude share, plain increments", dict(pd="crude_shift"), "rung", None),
    ("audit route: crude share of every loan, over 1 - C(a)", dict(pd="crude_r4", closed_only=False), "rung", None),
    ("audit route: crude share of every loan, over the share still open", dict(pd="crude_S", closed_only=False), "rung", None),
    ("audit route: the crude curve read from month on book 0", dict(pd="crude_uncond"), "rung", None),
    ("audit route: repayment counted again on EAD", dict(ead_runoff=True), "rung", None),
    ("risk grade added, held from the take-over", dict(by=("product", "risk_grade")), "c4", "SF7"),
    ("the stage 2 scalar on the default rate", dict(scalar_on="hazard"), "c4", "SF6"),
    ("v3 origination scores read as v4", dict(map_v3=False), "c4", "SF12"),
    ("hardship left out of staging", dict(use_hardship=False), "c4", "SF11"),
    ("base scenario only", dict(mult=1.0), "c4", "SF13"),
    ("EAD held at the current balance", dict(flat_ead=True), "c4", "SF14"),
    ("no discounting", dict(disc=False), "c4", "SF15"),
    ("the taken-over loans left out", dict(entry="own"), "same", "SF3"),
    ("each a and k read off the loans followed that far (landmark)", dict(pd="landmark"), "same", "SF8"),
    ("months on book from the first repayment", dict(clock="first"), "same", "SF5"),
    ("past-maturity loans carry nothing", dict(pm_zero=True), "same", "SF9"),
    ("hardship staged on the status label", dict(rule="label"), "same", "SF11"),
]


def main_call():
    rec = E.record(T)
    s = E.solve(T, rec=rec)
    A = s["total"]
    ok(abs(s["mult"] - 1.0555) < 1e-12, "SF13: the probability-weighted PD multiplier is 1.0555 (55 / 30 / 15 per cent on 1.00, "
       "1.28, 0.81)")
    for n in FIG:
        ok(edge(s[n], 1e5) >= 0.2, "%s sits %.2f of a bin from its rounding edge at AUD 0.1 million (%.0f)" % (n, edge(s[n], 1e5), s[n]))
    ok(round(M(s["stage1"]) + M(s["stage2"]), 1) == M(A), "SF16: the reported stage amounts add to the reported allowance")
    out = {}
    for name, kw, kind, sf in MAIN:
        x = E.solve(T, rec=rec, **kw)
        out[name] = x
        d = x["total"] / A - 1
        if kind in ("rung", "c4"):
            ok(M(x["total"]) != M(A), "%s%s lands off the allowance's bin (%.0f, %+.2f%%)" % (sf + ": " if sf else "", name, x["total"], 100 * d))
        else:
            ok(all(M(x[n]) == M(s[n]) for n in FIG), "%s: %s lands all three figures in their bins (%.0f, %+.2f%%)" % (sf, name, x["total"], 100 * d))
    r0, r1, r2 = (out[n]["total"] for n, _, k, _ in MAIN[:3])
    ok(r0 < r1 < r2 < A and A / r0 - 1 >= 0.10, "the ladder climbs: default path %.0f, grade %.0f, censored repayment %.0f, the "
       "answer %.0f (%.1f per cent above the default path)" % (r0, r1, r2, A, 100 * (A / r0 - 1)))
    ok(all(M(out[MAIN[0][0]][f]) != M(s[f]) for f in FIG), "the default path misses the allowance and both stage amounts")
    for n_ in (MAIN[1][0], MAIN[2][0]):
        ok(M(out[n_]["total"]) != M(A) and M(out[n_]["stage1"]) != M(s["stage1"]),
           "%s misses the allowance and stage 1 (stage 2 %.0f: %s)" % (n_, out[n_]["stage2"],
                                                                      "its bin" if M(out[n_]["stage2"]) == M(s["stage2"]) else "off"))
    ok(len({M(out[n]["total"]) for n, _, k, _ in MAIN if k == "rung"}) >= 7, "the rungs and audit routes land on distinct figures")
    same = [out[n]["total"] for n, _, k, _ in MAIN if k == "same"]
    ok(1.315e7 < min(same + [A]) and max(same + [A]) < 1.325e7, "C3 corridor: the answer and every convergent reading sit inside "
       "AUD 13.15 to 13.25 million (%.0f to %.0f)" % (min(same + [A]), max(same + [A])))
    ok(s["n3"] > 0 and s["n_pm"] >= 10 and s["ecl_pm"] > 0, "SF1: %d stage 3 loans sit outside the model; SF9: %d past-maturity "
       "loans carry an allowance (AUD %.0f)" % (s["n3"], s["n_pm"], s["ecl_pm"]))
    vfig, vst, _ = V.allowance(T)
    ok(all(abs(vfig[n] - s[n]) < 0.01 for n in FIG), "the independent verifier agrees to the cent")
    for how, name in (("drawdown", MAIN[0][0]), ("grade", MAIN[1][0]), ("km", MAIN[3][0]), ("own", "the taken-over loans left out"),
                      ("first", "months on book from the first repayment"), ("crude", MAIN[4][0]),
                      ("r4b", "audit route: repayment counted again on EAD")):
        vx, _, _ = V.allowance(T, how)
        ok(abs(vx["total"] - out[name]["total"]) < 0.01, "the verifier reproduces %s to the cent" % name)
    return s, out, rec


# ---------------------------------------------------------------------------------------------------------- the three rules

def three_rules(s, out, rec):
    """R1 the field's raw material ships; R2 every mess item is neutral to the call; R3a the easy route with every filed
    rule executed files the named wrong answer; R3b changing only the decisive step files the answer; R3c nothing planted"""
    tmp = tempfile.mkdtemp()
    try:
        cut = os.path.join(tmp, "norecord")
        shutil.copytree(T, cut)
        os.remove(os.path.join(cut, E.F["record"]))
        try:
            E.solve(cut)
            needs = False
        except Exception:
            needs = True
        ok(needs, "R1: the allowance cannot be computed without the loan record")
    finally:
        shutil.rmtree(tmp)
    cols = set()
    for f in os.listdir(T):
        if f.endswith(".csv"):
            cols |= set(pd.read_csv(os.path.join(T, f), nrows=1).columns)
        if f.endswith(".parquet"):
            cols |= set(pq.read_schema(os.path.join(T, f)).names)
    ok(not any(re.search(r"\bpd\b|hazard|month_on_book|\bmob\b|stage|lgd|loss|survival|entry|migrat", c.lower()) for c in cols),
       "R1: no shipped column carries a PD, a rate, a month on book, a stage or a loss")
    hz = E.hazards(rec, "transfer")
    ok(all(hz[p][2][1] > 30000 for p in ("PL36", "PL60")), "R1: the PD is estimated from tens of thousands of loans per product")
    r2 = {"the register's Tuesday status batch (staging on the label)": out["hardship staged on the status label"],
          "months on book from dates (counted from the first repayment)": out["months on book from the first repayment"],
          "past-maturity loans (read as carrying nothing)": out["past-maturity loans carry nothing"]}
    bk = pd.read_csv(os.path.join(T, E.F["book"]))
    for name, x in r2.items():
        ok(all(M(x[n]) == M(s[n]) for n in FIG), "R2: %s keeps the call" % name)
    tmp = tempfile.mkdtemp()
    try:
        cut = os.path.join(tmp, "nostate")
        shutil.copytree(T, cut)
        b2 = pd.read_csv(os.path.join(cut, E.F["book"]))
        b2["state"] = "NSW"
        b2["channel"] = "Direct"
        b2.to_csv(os.path.join(cut, E.F["book"]), index=False)
        ok(abs(E.solve(cut)["total"] - s["total"]) < 0.01, "R2: the blank states of overseas addresses, and state and channel "
           "altogether, move nothing")
    finally:
        shutil.rmtree(tmp)
    hd = pd.read_excel(os.path.join(T, E.F["hardship"]), sheet_name=0)
    covid = hd[(hd["hardship_reason"] == "COVID-19") & (hd["arrangement_status"] == "Active")]
    ok(len(covid) >= 300 and not covid["loan_no"].isin(set(bk["loan_no"])).any(), "R2: the bulk-loaded COVID-19 rows still "
       "reading Active (%d) sit on loans closed before the reporting date" % len(covid))
    easy = out[MAIN[0][0]]
    ok(M(easy["total"]) != M(s["total"]) and abs(M(easy["total"]) - M(s["total"])) >= 1.0 and easy["n2"] == s["n2"]
       and easy["mult"] == s["mult"], "R3a: the easy route with every filed rule executed files AUD %.1f million against %.1f, "
       "on the same stages and multiplier" % (M(easy["total"]), M(s["total"])))
    one = E.solve(T, rec=rec, entry="transfer")
    ok(abs(one["total"] - s["total"]) < 0.01, "R3b: changing only the decisive step files the answer")
    for name in (MAIN[1][0], MAIN[2][0]):
        ok(M(out[name]["total"]) != M(s["total"]), "R3b: %s alone leaves the easy route off the answer" % name)
    names = sorted(os.listdir(T))
    ok(not any(re.search(r"memo|draft|note_to|briefing|prior|q2|ecl_allowance", f.lower()) for f in names),
       "R3c: no memo, draft, briefing or prior allowance in the pack")
    pol = docx_text(POLICY())
    dic = open(os.path.join(T, "ecl_extract_dictionary.txt")).read()
    for v in ("13.2", "10.6", "2.6", "13,178", "13178", "11.6", "11,551"):
        ok(re.search(r"(?<![\d.])%s(?![\d])" % re.escape(v), pol + dic) is None, "R3c: no document states %s" % v)
    for f in (E.F["book"], E.F["june"]):
        x = pd.read_csv(os.path.join(T, f)).drop(columns=["loan_no"])
        ok(x.duplicated().sum() <= 0.001 * len(x), "R3c: %s repeats no block of rows (%d pairs of loans on identical terms "
           "and balances)" % (f, x.duplicated().sum()))
    r = pd.read_parquet(os.path.join(T, E.F["record"])).drop(columns=["loan_no"])
    ok(r.duplicated().sum() < 0.002 * len(r), "R3c: the record repeats no block of rows (%d coincident rows)" % r.duplicated().sum())
    return dict(easy=easy["total"], one=one["total"])


# ---------------------------------------------------------------------------------------------------------- forks on the call

def forks():
    """SF rows of the main call that rest on a filed sentence"""
    pol = docx_text(POLICY())
    ok("in stages 1 and 2. Loans in stage 3 (90 or more days past due) are assessed under CRM-ECL-06 and are outside this method."
       in pol, "SF1: 1.1 puts stage 3 outside the collective allowance")
    ok("from the default and repayment experience of every loan of its product in the loan record" in pol
       and "it is the probability that a loan of the product that is open and not in default at the end of month on book a "
           "first defaults in month on book a + k" in pol, "SF3, SF4: 3.2 files the estimand: every loan of the product in the "
                                                          "record, open and not in default at a, first defaulting at a + k")
    ok("Months on book are calendar months counted from the month the loan was drawn, so a loan is in month on book 0 in the "
       "month it is drawn." in pol, "SF5: 3.2 counts months on book from the drawdown month")
    # SF35: exit at first default (round 2c re-judge's residual fork). 3.2's estimand is the probability that a loan open
    # and not in default at a first defaults at a + k, so a loan leaves the loans at risk at its first default. Keeping
    # the defaulted loans that were not written off or settled at risk from their default month (an instant cure) books
    # AUD 13.1 million; the record carries no cure date, and 343 of the 761 still open are 90 or more days past due at 30
    # September, so the data do not support an instant cure either
    rec_ = E.record(T)
    kept = (rec_["ev"] == "D") & ~rec_["closure_reason"].isin(["Written off", "Settled"])
    cd_ = pd.to_datetime(rec_["closed_date"])
    cm_ = (cd_.dt.year * 12 + cd_.dt.month - 1)
    rec_["xr"] = np.where(kept, np.where(cm_.notna(), cm_.fillna(0).astype(int), E.RM) - rec_["draw_m"], rec_["x"])
    ic = E.solve(T, rec=rec_)
    bk_ = pd.read_csv(os.path.join(T, E.F["book"]))
    op = rec_[kept & (rec_["status"] == "Open")].merge(bk_[["loan_no", "days_past_due"]], on="loan_no")
    ok("first defaults in month on book a + k" in pol and "default_month" in rec_.columns and not any("cure" in c for c in rec_.columns)
       and int((op["days_past_due"] >= 90).sum()) >= 0.3 * len(op) and M(ic["total"]) == 13.1,
       "SF35: kept at risk from the default month, the %d defaulted loans not written off or settled book AUD %.1f million "
       "(%.1f, %.1f); ruled out by 3.2's 'first defaults' and by the record (no cure date; %d of the %d still open are 90 or "
       "more days past due at 30 September)" % (int(kept.sum()), M(ic["total"]), M(ic["stage1"]), M(ic["stage2"]),
                                               int((op["days_past_due"] >= 90).sum()), len(op)))
    ok("Stage 2 loans carry the significant-increase scalar of 2.4 on PD(k)." in pol, "SF6: the scalar sits on PD(k)")
    ok("PD is set by product and month on book only. Risk grade, channel and state do not enter it, and the record's experience "
       "is used as it stands, with no curve fitted to it or smoothing." in pol, "SF7, SF8: 7.1 sets PD by product and month on "
                                                                                "book, the experience as it stands")
    ok("Repayment dates during a payment pause count as repayments fallen due" in pol
       and "A loan past its contractual maturity that is not in stage 3 has a remaining term of one month." in pol,
       "SF9, SF10: 3.2 pins pause dates and past-maturity loans")
    ok("A hardship arrangement is in force from its start date to its end date; one that is cancelled stops being in force "
       "on the day it is cancelled." in pol, "SF11, SF28: 2.2 states the in-force rule and when a cancelled arrangement stops")
    ok("Origination scores recorded on the v3 scorecard are converted to v4 with the published migration formula" in pol,
       "SF12: 2.3 converts v3")
    ok("scales PD(k) by its multiplier. The allowance is the probability-weighted sum of the scenario ECLs." in pol,
       "SF13: section 4 weights the scenarios")
    ok("repaid over the remaining term in equal monthly instalments at the contract rate" in pol
       and "(1 + contract rate / 12) to the power -k" in pol, "SF14, SF15: 3.4 and 3.5 fix EAD and DF")
    ok("figures are rounded only when they are reported" in pol, "SF16: 3.7 rounds only when reported")
    for w in ("delayed", "truncat", "entry", "survivor", "taken over", "take-over", "go-live", "earlier platform", "immortal",
              "risk set", "kaplan", "hazard", "life table", "aalen", "at risk"):
        ok(w not in pol.lower(), "no filed sentence in CRM-ECL-04 names the method or the take-over (%r)" % w)
    dic = " ".join(open(os.path.join(T, "ecl_extract_dictionary.txt")).read().split())
    ok("when it took over the personal loan book from the earlier platform" in dic, "the dictionary files the record's population")
    for w in ("delayed", "truncat", "late entry", "entry date", "survivor", "risk set", "hazard", "count from", "months held",
              "immortal", "at risk", "only from"):
        ok(w not in dic.lower(), "the dictionary names no method (%r)" % w)
    rec = pd.read_parquet(os.path.join(T, E.F["record"]))
    ok(rec["default_month"].dropna().max() <= "2026-09" and (rec["closed_date"].dropna().astype(str) <= "2026-09-30").all(),
       "SF17, SF18: the record runs to 30 September 2026, defaults at the month end 1.2 and the loan system record")


# ---------------------------------------------------------------------------------------------------------- realism

def realism():
    bk = pd.read_csv(os.path.join(T, E.F["book"]))
    first = bk["first_repayment_date"].map(E._d)
    draw = bk["drawdown_date"].map(E._d)
    dpd = bk["days_past_due"].values
    nomatch = gt = before = 0
    for f, d, t in zip(first, dpd, bk["term_months"]):
        if d == 0:
            continue
        sched = due_dates(f, int(t) + 6)
        nomatch += (R - dt.timedelta(days=int(d))) not in sched
        before += sched[0] >= R
        gt += d > (R - sched[0]).days
    ok(nomatch == 0 and gt == 0 and before == 0, "days past due match a due date (%d), never exceed days since the first (%d), none before it (%d)" % (nomatch, gt, before))
    ok(dpd.max() <= 179, "no loan on the book at 180 or more days past due (%d)" % dpd.max())
    a = np.array([E.months_due(f) for f in first])
    i = bk["contract_rate"].values / 12
    n = bk["term_months"].values.astype(float)
    amt = bk["original_amount"].values.astype(float)
    pay = amt * i / (1 - (1 + i) ** -n)
    k = np.minimum(a, n)
    sched = amt * (1 + i) ** k - pay * ((1 + i) ** k - 1) / i
    on = (np.abs(bk["current_balance"].values - np.round(sched, 2)) < 0.005) & (a >= 1)
    ok(on.mean() < 0.01, "balances on the textbook schedule to the cent are under 1 per cent of repaying loans (%.4f)" % on.mean())
    arr = (dpd >= 30) & (a < n)
    ok((bk["current_balance"].values[arr] > sched[arr]).mean() > 0.8, "loans 30 or more days past due sit above schedule")
    hd = pd.read_excel(os.path.join(T, E.F["hardship"]), sheet_name=0)
    st, en = hd["arrangement_start"].dt.date, hd["arrangement_end"].dt.date
    ok(all(is_bd(d) for d in st), "arrangements start on business days")
    life = dict(zip(bk["loan_no"], zip(draw, [due_dates(f, int(t))[-1] for f, t in zip(first, bk["term_months"])])))
    inside = [(life[l][0] < s0 and e0 <= life[l][1]) for l, s0, e0 in zip(hd["loan_no"], st, en) if l in life]
    ok(all(inside), "arrangements on book loans fall inside the loan's life")
    ok(((hd["cancelled_date"].notna()) == (hd["arrangement_status"] == "Cancelled")).all(), "cancelled_date is filled exactly for cancelled arrangements")
    ok(((en - st).map(lambda x: x.days) >= 1).all(), "every arrangement ends after the day it starts")
    cx = hd[hd["arrangement_status"] == "Cancelled"]
    ok((cx["cancelled_date"] <= cx["arrangement_end"]).all(), "arrangements are cancelled inside their term")
    ok((hd["received_date"].diff().dropna().dt.days >= 0).all() and hd["case_ref"].is_monotonic_increasing, "case numbers run in order of receipt")
    ok(hd.loc[hd["arrangement_status"] == "Active", "arrangement_type"].nunique() == 2, "active arrangements are of both types")
    ok((hd["loan_no"].value_counts() > 1).sum() > 100, "some loans have repeat arrangements")
    ok(hd["received_date"].min().date() >= W.GO_LIVE and hd["received_date"].min().date() < dt.date(2020, 3, 1),
       "the register runs from soon after the go-live, nothing from the earlier platform")
    dur = [(e0 - s0).days + 1 for s0, e0 in zip(st, en)]
    ok(np.mean([d in (60, 90, 180) for d in dur]) < 0.15, "durations run in calendar months, not fixed day counts")
    act = hd[(hd["arrangement_status"] != "Cancelled") & (st <= R) & (en >= R)]
    ok((en != R).all() and (hd["cancelled_date"].dropna().dt.date < R).all(), "no arrangement ends, or is cancelled, on the reporting date")
    ok((hd["cancelled_date"].dropna().dt.date != JUNE).all(), "no arrangement is cancelled on 30 June either")
    ok(len(set(act["loan_no"]) & set(bk["loan_no"])) == len(set(act["loan_no"])), "every arrangement in force belongs to a book loan")
    wo = pd.read_csv(os.path.join(T, "write_offs_2025-07_2026-09.csv"))
    bands = [((dpd >= x) & (dpd < x + 30)).sum() for x in (90, 120, 150)]
    ok(0.6 <= wo["loans_written_off"].mean() / np.mean(bands) <= 1.4, "monthly write-offs agree with the stage 3 bands (%.1f vs %.1f)" % (wo["loans_written_off"].mean(), np.mean(bands)))
    ok(all(is_bd(x) for x in draw), "loans are funded on NSW business days")
    ok((draw.map(lambda x: x.day) >= 29).sum() > 100, "funding on the 29th to 31st")
    gap = np.array([(f - x).days for f, x in zip(first, draw)])
    ok(gap.min() >= 14 and gap.max() <= 44, "first repayment 2 to 6 weeks after funding")
    same = np.mean([f == dt.date(x.year + (x.month == 12), x.month % 12 + 1, min(x.day, 28)) for f, x in zip(first, draw)])
    ok(0.2 < same < 0.6, "about four in ten first repayments one month after funding")
    rec = pd.read_parquet(os.path.join(T, E.F["record"]))
    bm = pd.to_datetime(rec["drawdown_date"]).dt.strftime("%Y-%m").value_counts()
    ok(bm["2020-04"] < 0.7 * bm["2020-02"] and bm["2025-12"] < bm["2025-11"], "volumes carry the 2020 shock and December seasonality")
    f3, f4 = dict(W.BANDS_V3), dict(W.BANDS_V4)
    band = [next((g for g in W.GRADES if sc >= (f4 if card == "v4" else f3)[g]), "E")
            for sc, card in zip(bk["origination_score"], bk["origination_scorecard"])]
    ok(np.mean(np.array(band) == bk["risk_grade"].values) > 0.97, "grades follow the score bands")
    for col in ("origination_score", "current_score_v4"):
        vc = bk[col].value_counts()
        ok(vc.get(bk[col].max(), 0) < 5, "%s has no pile-up at its maximum" % col)
    med = bk.groupby(["risk_grade", "product"])["contract_rate"].median().unstack()
    ok((med["PL60"] > med["PL36"]).sum() >= 4, "PL60 priced above PL36 in most grades")
    yr = draw.map(lambda x: x.year)
    c = bk[(bk["risk_grade"] == "C") & (bk["product"] == "PL60")].assign(y=yr)
    ok(c[c.y == 2024]["contract_rate"].median() - c[c.y == 2022]["contract_rate"].median() > 0.015, "rates follow the cash rate cycle")
    rr = rec.assign(y=pd.to_datetime(rec["drawdown_date"]).dt.year)
    rc = rr[(rr["risk_grade"] == "C") & (rr["product"] == "PL60")].groupby("y")["contract_rate"].median()
    ok(rc[2017] > rc[2021] and rc[2024] > rc[2021], "the earlier platform's loans were priced on the higher cash rates before 2020")
    vc = bk["original_amount"].value_counts()
    ok(vc.index[0] < 50000 and bk["original_amount"].max() > 50000, "the cap is not the most common amount")
    ch = bk["channel"].value_counts(normalize=True)
    ok(ch.max() - ch.min() > 0.2 and abs(bk["product"].eq("PL36").mean() - 0.5) > 0.01, "channel and product mixes are not even splits")
    vs = bk["state"].value_counts()
    shp = 100 * vs / len(bk)
    ok(vs.is_unique and (np.abs(shp - np.round(shp)) >= 0.15).sum() >= 6, "state counts all differ and sit off whole percentages")
    o4 = np.where(bk["origination_scorecard"] == "v3", np.floor(1.12 * bk["origination_score"] - 18 + 0.5), bk["origination_score"])
    fresh = a == 0
    cur = bk["current_score_v4"].values
    ok(fresh.sum() > 500 and (cur[fresh] == o4[fresh]).all(), "behaviour score is the v4 origination score until the first repayment falls due (%d loans)" % fresh.sum())
    ok((cur[~fresh] != o4[~fresh]).mean() > 0.9, "behaviour scores move once repayments fall due")
    num = bk["loan_no"].str[3:9].astype(int)
    v3 = (bk["origination_scorecard"] == "v3").values
    sw = dt.date(2024, 7, 1)
    ok(num[v3].max() < num[~v3].min(), "every v3 loan number precedes every v4 loan number")
    ok((draw[v3] >= sw).sum() >= 50 and (draw[~v3] >= sw).all(), "loans applied for before 1 July 2024 and funded after it stay on v3")
    luhn = lambda q: (10 - sum((lambda x: x - 9 if x > 9 else x)(int(c) * (2 if i % 2 == 0 else 1)) for i, c in enumerate(q[::-1])) % 10) % 10
    ok(all(int(x[-1]) == luhn(x[3:9]) for x in rec["loan_no"]), "loan numbers carry a Luhn check digit, the earlier platform's too")
    pz = act.loc[act["arrangement_type"] == "Payment pause", "loan_no"]
    dp = bk.set_index("loan_no").loc[pz, "days_past_due"]
    ok(len(pz) > 50 and (dp == 0).all(), "no arrears on a loan under a payment pause in force on the reporting date (%d pauses)" % len(pz))
    red = hd.loc[hd["arrangement_type"] == "Reduced repayments", "agreed_repayment_aud"]
    ok((red % 5 == 0).mean() < 0.75 and (red % 1 != 0).mean() > 0.1 and (red <= 50).mean() > 0.03, "agreed repayments are a realistic mix")
    rec_ = wo["recoveries_aud"].sum() / wo["gross_write_off_aud"].sum()
    ok(0.15 <= rec_ <= 0.35, "recoveries run at %.3f of gross write-offs over the 15 months, the March 2026 sale included" % rec_)
    ok(wo["loans_written_off"].mean() < 0.95 * ((dpd >= 90) & (dpd < 120)).sum(), "fewer loans are written off each month than reach 90 days past due")


# ---------------------------------------------------------------------------------------------------------- the June extract

def june_extract():
    bk = pd.read_csv(os.path.join(T, E.F["book"]))
    jb = pd.read_csv(os.path.join(T, E.F["june"]))
    ok(list(jb.columns) == [c for c in bk.columns if c != "current_score_v4"], "the June extract has the September columns but the score")
    ok(jb["loan_no"].is_unique and jb["loan_no"].is_monotonic_increasing, "the June extract is one row per loan, by loan number")
    both = jb.merge(bk, on="loan_no", suffixes=("_j", "_s"))
    for c in ("product", "risk_grade", "drawdown_date", "first_repayment_date", "term_months", "contract_rate", "original_amount",
              "origination_score", "origination_scorecard", "channel"):
        ok((both[c + "_j"] == both[c + "_s"]).all(), "%s agrees across the two extracts" % c)
    ok(((both["state_j"] == both["state_s"]) | (both["state_j"].isna() & both["state_s"].isna())).all(), "state agrees across the two extracts")
    new = bk[~bk["loan_no"].isin(set(jb["loan_no"]))]
    ok((pd.to_datetime(new["drawdown_date"]) > pd.Timestamp(JUNE)).all(), "every September loan missing from June was drawn after 30 June (%d)" % len(new))
    up = both[(both["days_past_due_j"] == 0) & (both["days_past_due_s"] == 0)]
    ok((up["current_balance_s"] < up["current_balance_j"]).mean() > 0.95, "up-to-date loans pay their balances down over the quarter")
    # B3: the June extract counts days past due to its run date, 1 July, as the dictionary files. The records agree
    # three ways: every loan in arrears sits that many days after a due date counted to 1 July, not 30 June;
    # the month-end batch writes off at 180 days at close of business (1.4), yet loans show 180; and the loan record dates
    # the default of every loan the extract shows at exactly 90 days to a later month end than June's
    run = JUNE + dt.timedelta(days=1)
    late = jb[jb["days_past_due"] > 0]
    to_run = to_close = 0
    for f, d, t in zip(late["first_repayment_date"].map(E._d), late["days_past_due"], late["term_months"]):
        sched = set(due_dates(f, int(t) + 6))
        to_run += (run - dt.timedelta(days=int(d))) in sched
        to_close += (JUNE - dt.timedelta(days=int(d))) in sched
    ok(to_run == len(late) and to_close <= 0.5 * len(late), "SF25: June days past due count to the 07:30 run on 1 July: %d of %d "
       "sit on a due date counted to 1 July, %d counted to 30 June" % (to_run, len(late), to_close))
    at180 = int((jb["days_past_due"] == 180).sum())
    ok(jb["days_past_due"].max() == 180 and at180 >= 5, "SF25: the June batch wrote off at 180 days; %d loans at 179 then read 180 at the run" % at180)
    rec = pd.read_parquet(os.path.join(T, E.F["record"]), columns=["loan_no", "default_month"])
    x = jb.merge(rec, on="loan_no", how="left")
    at90 = x[x["days_past_due"] == 90]
    over = x[x["days_past_due"] > 90]
    ok(len(at90) >= 3 and (at90["default_month"] > "2026-06").all() and (over["default_month"] <= "2026-06").all(),
       "SF25: the loan record dates the default of each of the %d loans the June extract shows at 90 days to a later month end, "
       "and of every loan above 90 to June or before (%d)" % (len(at90), len(over)))
    jc = (jb["days_past_due"] - 1).clip(lower=0)
    r_j, r_s = (jc >= 30).mean(), (bk["days_past_due"] >= 30).mean()
    ok(0.75 <= r_j / r_s <= 1.33, "arrears of 30 days or more run at a like share in June and September (%.4f, %.4f)" % (r_j, r_s))
    closed = jb[~jb["loan_no"].isin(set(bk["loan_no"]))]
    rec2 = pd.read_parquet(os.path.join(T, E.F["record"]), columns=["loan_no", "status", "closed_date"])
    cj = closed.merge(rec2, on="loan_no")
    ok(len(cj) == len(closed) and (cj["status"] == "Closed").all() and (cj["closed_date"].astype(str) > "2026-06-30").all(),
       "SF27: every June loan missing from September closed in the quarter, as the record shows (%d)" % len(closed))
    return len(jb), len(closed), len(new), len(late), to_close, at180, len(at90)


# ---------------------------------------------------------------------------------------------------------- the scoring runs

def scoring_runs():
    sc = E.score_runs(T)
    ok(list(sc.columns) == ["loan_no", "run_date", "score_v4", "score_status"], "the scoring file carries loan, run date, "
                                                                               "score and status")
    runs = sorted(sc["run_date"].unique())
    ok(runs == sorted([dt.date(2026, m, 1) for m in range(4, 10)] + [RERUN]), "six month-end runs made on the first day of "
       "April to September 2026, and the June month end's partition made again on 2 July")
    ok(not sc.duplicated(["loan_no", "run_date"]).any(), "one score per loan per run")
    # SF34: the June month end's run. The 1 July run could not score one partition of the engine's book and wrote each of
    # those loans its last score as Carried forward; the partition was run again on 2 July. Every loan scored for the June
    # month end has exactly one Scored row among the rows made after 30 June and before the July month end's run
    cf = sc[sc["score_status"] == "Carried forward"]
    ok(set(sc["score_status"]) == {"Scored", "Carried forward"} and (cf["run_date"] == dt.date(2026, 7, 1)).all(),
       "SF34: Carried forward appears only in the run dated 1 July (%d loans)" % len(cf))
    rr = sc[sc["run_date"] == RERUN]
    ok(set(rr["loan_no"]) == set(cf["loan_no"]) and (rr["score_status"] == "Scored").all(), "SF34: the 2 July rows score "
       "exactly the loans carried forward on 1 July")
    jm = sc[(sc["run_date"] > JUNE) & (sc["run_date"] < dt.date(2026, 8, 1)) & (sc["score_status"] == "Scored")]
    ok(jm["loan_no"].is_unique, "SF34: each loan has one Scored row for the June month end")
    may = sc[sc["run_date"] == dt.date(2026, 6, 1)].set_index("loan_no")["score_v4"]
    has_may = cf["loan_no"].isin(may.index)
    ok((cf.loc[has_may, "score_v4"].to_numpy() == may.loc[cf.loc[has_may, "loan_no"]].to_numpy()).all() and has_may.mean() > 0.9,
       "SF34: each carried-forward score repeats the loan's May month-end score (%d of %d; the rest first scored in June "
       "carry the application score)" % (int(has_may.sum()), len(cf)))
    alt = rr.set_index("loan_no")["score_v4"]
    ok((alt.loc[cf["loan_no"]].to_numpy() != cf["score_v4"].to_numpy()).mean() > 0.8, "SF34: the 2 July score differs from "
       "the carried-forward one for most of the partition")
    meta = meta_of(os.path.join(T, E.F["scores"]))
    ok(meta["runs"] == "the March, April, May, June, July and August 2026 month ends" and meta["run_date"] == "date the run was made"
       and all(c in meta for c in sc.columns), "SF26: the file names the six month ends it scores, dates each run by the day "
                                               "it was made, and describes every column in one form")
    jb = pd.read_csv(os.path.join(T, E.F["june"]))
    jun = jm
    due = set(jb.loc[pd.to_datetime(jb["first_repayment_date"]) <= pd.Timestamp(JUNE), "loan_no"])
    ok(set(jun["loan_no"]) == due, "the June month-end run scores exactly the June loans whose first repayment has fallen due (%d)" % len(due))
    bk = pd.read_csv(os.path.join(T, E.F["book"]))
    orig = pd.concat([jb, bk]).drop_duplicates("loan_no").set_index("loan_no")
    o4 = np.where(orig["origination_scorecard"] == "v3", np.floor(1.12 * orig["origination_score"] - 18 + 0.5), orig["origination_score"])
    o4 = pd.Series(o4, index=orig.index)
    shares = {}
    for d in runs:
        if d == RERUN:
            continue
        x = sc[(sc["run_date"] == d) & (sc["score_status"] == "Scored")] if d != dt.date(2026, 7, 1) else jm
        x = x[x["loan_no"].isin(o4.index)]
        shares[d.isoformat()] = float(((o4.loc[x["loan_no"]].to_numpy() - x["score_v4"].to_numpy()) >= 60).mean())
    vals = list(shares.values())
    ok(0.05 <= min(vals) and max(vals) <= 0.09 and max(vals) - min(vals) <= 0.015,
       "the share 60 points or more below origination holds steady across the runs (%s)" % {k: round(v, 4) for k, v in shares.items()})
    jl = sc[sc["run_date"] == dt.date(2026, 8, 1)].set_index("loan_no")["score_v4"]
    jn = jun.set_index("loan_no")["score_v4"]
    chg = (jl - jn).dropna().abs()
    ok(3 <= chg.median() <= 15, "a typical score moves %.0f points from one month-end run to the next" % chg.median())
    aug = sc[sc["run_date"] == dt.date(2026, 9, 1)].set_index("loan_no")["score_v4"]
    sep = bk.set_index("loan_no")["current_score_v4"]
    common = aug.index.intersection(sep.index)
    ok((aug.loc[common] != sep.loc[common]).mean() > 0.7, "the September run in the extract is not a copy of the August run")
    return shares


# ---------------------------------------------------------------------------------------------------------- Collections files

def collections_files():
    fy, led, sale, price, head = E.collections(T)
    ok(fy["loan_no"].is_unique, "SF19: one row per 2023-24 default")
    ok(fy["default_month"].between("2023-07", "2024-06").all(), "SF19, SF20: every default month falls in July 2023 to June 2024")
    ok(set(fy["resolution"]) == {"Cured", "Settled", "Written off"}, "defaults are cured, settled or written off")
    dme = pd.to_datetime(fy["default_month"] + "-01") + pd.offsets.MonthEnd(0)
    ok((pd.to_datetime(fy["resolution_date"]) > dme).all(), "every resolution falls after the default month end")
    wo = fy[fy["resolution"] == "Written off"]
    gap = (pd.to_datetime(wo["resolution_date"]).dt.to_period("M") - pd.to_datetime(wo["default_month"]).dt.to_period("M")).map(lambda x: x.n)
    ok((pd.to_datetime(wo["resolution_date"]) == pd.to_datetime(wo["resolution_date"]) + pd.offsets.MonthEnd(0)).all()
       and gap.between(2, 4).all(), "write-offs fall in a month-end batch two to four months after the default month")
    q = [((int(m[5:7]) - 7) % 12) // 3 + 1 for m in fy["default_month"]]
    cells = Counter(zip(fy["product"], q))
    ok(min(cells.values()) >= 100, "every product and quarter holds 100 or more defaults (%s)" % dict(cells))
    rec = pd.read_parquet(os.path.join(T, E.F["record"]))
    rj = fy.merge(rec[["loan_no", "default_month", "closure_reason"]], on="loan_no", suffixes=("", "_r"))
    ok(len(rj) == len(fy) and (rj["default_month"] == rj["default_month_r"]).all(), "the record dates every 2023-24 default to the "
       "month the Collections MI does")
    ok(led["receipt_ref"].is_unique and (led["receipt_ref"].str[2:].astype(int).diff().dropna() > 0).all()
       and pd.to_datetime(led["received_date"]).is_monotonic_increasing, "receipt numbers run in date order")
    ok(led["received_date"].between("2023-07-01", "2026-09-30").all(), "the receipts file runs July 2023 to September 2026")
    blank = led[led["loan_no"].isna()]
    ok(len(blank) == 1 and blank["source"].iloc[0] == "Debt sale", "the one receipt with no loan is the sale's")
    amt = led["amount_aud"].astype(float)
    neg = led[amt < 0]
    ok(len(neg) >= 5 and (neg["source"] == "Direct debit").all() and neg["narrative"].str.startswith("Dishonour of RC").all(),
       "dishonoured direct debits are reversed by negative entries that name the receipt (%d)" % len(neg))
    by_ref = led.set_index("receipt_ref")
    for _, r in neg.iterrows():
        o = by_ref.loc[r["narrative"].split()[-1]]
        ok(o["loan_no"] == r["loan_no"] and abs(float(o["amount_aud"]) + float(r["amount_aud"])) < 0.005
           and o["received_date"][:7] == r["received_date"][:7] and o["received_date"] < r["received_date"],
           "reversal %s nets its receipt inside the month" % r["receipt_ref"])
    st = led[led["source"] == "Settlement"]
    ok((st["amount_aud"].astype(float) % 50 == 0).all(), "settlements are agreed in round fifty dollars")
    # A1: Collections closes a default case as cured when it approves a collections arrangement, dated from the
    # arrangement's start, and does not reopen it when the arrangement is withdrawn; the register keeps the arrangement as
    # cancelled before its start, and the record shows how each loan ended
    withdrawn, stale = E.withdrawn_cures(T, fy, led)
    caught = withdrawn - stale
    hd = pd.read_excel(os.path.join(T, E.F["hardship"]), sheet_name=0)
    ok(len(stale) >= 30 and len(caught) >= 8, "SF22: %d cases closed as cured on a withdrawn arrangement were then written off or "
       "settled, %d caught up instead" % (len(stale), len(caught)))
    cq = fy.assign(q=q)
    for (p_, q_), g_ in cq.groupby(["product", "q"]):
        ok(g_["loan_no"].isin(stale).sum() >= 2, "%s Q%d holds cases closed as cured on withdrawn arrangements and later written off" % (p_, q_))
    asof = E.attribute(T, led, "asof")
    rc = asof[asof["loan_no"].isin(withdrawn)].merge(fy[["loan_no", "default_month", "resolution_date"]], on="loan_no")
    rc = rc[(pd.to_datetime(rc["received_date"]) > pd.to_datetime(rc["default_month"] + "-01") + pd.offsets.MonthEnd(0))
            & (rc["received_date"] <= rc["resolution_date"])]
    big = rc.assign(a=rc["amount_aud"].astype(float)).groupby("loan_no")["a"].max()
    ok(set(big[big >= 500].index) == caught, "SF22: every caught-up case has its catch-up receipt and no case later written off has one")
    wd = hd[hd["loan_no"].isin(withdrawn) & (hd["arrangement_status"] == "Cancelled") & (hd["cancelled_date"] < hd["arrangement_start"])]
    ok(wd["loan_no"].nunique() == len(withdrawn) and (wd["arrangement_type"] == "Reduced repayments").all(),
       "each sits in the register as reduced repayments cancelled before its start")
    dm = pd.to_datetime(fy["default_month"] + "-01") + pd.offsets.MonthEnd(0)
    post = hd.merge(fy.assign(dme=dm)[["loan_no", "dme", "resolution", "resolution_date"]], on="loan_no")
    post = post[post["received_date"] > post["dme"]]
    ok(len(post) >= 100 and (post["resolution"] == "Cured").all()
       and (post["arrangement_start"] == pd.to_datetime(post["resolution_date"])).all(),
       "every arrangement approved after a 2023-24 default closed its case as cured, dated from its start (%d)" % len(post))
    cut, settled = W.SALE_CUTOFF, W.SALE_SETTLED
    ok(abs(round(float(sale["face_value_aud"].astype(float).sum()) * price, 2) - float(blank["amount_aud"].iloc[0])) < 0.01
       and abs(float(head["Consideration (AUD)"]) - float(blank["amount_aud"].iloc[0])) < 0.005
       and blank["received_date"].iloc[0] == settled.isoformat(),
       "SF24: the sale receipt is face value times price, as the schedule's header states, on the settlement date")
    ok(sale["loan_no"].is_unique and (sale["face_value_aud"].astype(float) > 0).all()
       and (sale["written_off"] <= W.SALE_WRITTEN_OFF_BY.strftime("%Y-%m")).all(), "sold accounts were written off by August 2025")
    fs = fy[fy["loan_no"].isin(set(sale["loan_no"]))]
    fq = Counter(zip(fs["product"], [((int(m[5:7]) - 7) % 12) // 3 + 1 for m in fs["default_month"]]))
    rs = fs.merge(rec[["loan_no", "closure_reason"]], on="loan_no")
    ok(len(fq) == 8 and min(fq.values()) >= 20 and (rs["closure_reason"] == "Written off").all()
       and set(rs.loc[rs["resolution"] != "Written off", "loan_no"]) <= stale,
       "the sale holds written-off 2023-24 defaults from every quarter and both products (%s), %d of them closed as cured on a "
       "withdrawn arrangement" % (dict(fq), int((rs["resolution"] != "Written off").sum())))
    # A2: Wirraway's placements. Every written-off account is placed with the agency after the write-off; from November 2024
    # the agency reissues the references of closed placements, and the receipts file posts each remittance to the loan of
    # the first placement that carried the reference it quotes
    pl = E.placements(T)
    ok(list(pl.columns[:6]) == ["Agency Ref", "Client Ref", "Placed", "Placed Balance", "Closed", "Close Reason"],
       "the placement report keeps the agency's own column names")
    ok(pl["Placed"].str.match(r"\d\d/\d\d/20\d\d$").all(), "the agency dates its report day first")
    ok(set(pl["Close Reason"].dropna()) == {"Paid in full", "Settled", "Returned", "Recalled"}, "placements close four ways")
    wo_r = rec[rec["closure_reason"] == "Written off"]
    placed = pl.merge(wo_r[["loan_no", "closed_date"]], left_on="Client Ref", right_on="loan_no", how="left")
    ok(placed["closed_date"].notna().all() and placed["Client Ref"].is_unique
       and ((placed["placed"] - pd.to_datetime(placed["closed_date"])).dt.days.between(1, 31)).all(),
       "every placement is a written-off loan, placed once, within a month of its write-off")
    many = pl["Agency Ref"].value_counts()
    multi = pl[pl["Agency Ref"].isin(many[many > 1].index)].sort_values("placed")
    second = multi.groupby("Agency Ref")["placed"].apply(lambda v: v.iloc[1])
    ok((many > 1).sum() >= 300 and second.min() >= pd.Timestamp(W.AGENCY_REISSUE_FROM), "SF23: %d references carried by more than "
       "one placement, none reissued before November 2024" % (many > 1).sum())
    clash = 0
    for ref, g in multi.groupby("Agency Ref"):
        c0 = g["closed"].iloc[:-1].to_numpy()
        p1 = g["placed"].iloc[1:].to_numpy()
        clash += int((pd.isna(c0) | ((p1 - c0) < np.timedelta64(90, "D"))).sum())
    ok(clash == 0, "SF23: a reference is reissued only after its placement closed, and at least 90 days after")
    ag = led[led["source"] == "Agency remittance"]
    ok(ag["narrative"].str.match(r"Wirraway Mercantile Agents remittance [A-Z][a-z]{2} 20\d\d ref \d{6}$").all(),
       "every agency remittance quotes the agency's reference")
    moved = (asof["loan_no"] != led["loan_no"]) & (led["source"] == "Agency remittance")
    ok(int(moved.sum()) >= 500, "SF23: %d remittances are posted to another loan than the placement that carried their reference "
       "on the day" % int(moved.sum()))
    post_sold = ag[ag["loan_no"].isin(set(sale["loan_no"])) & (ag["received_date"] > settled.isoformat())]
    ok(len(post_sold) >= 20 and asof.loc[post_sold.index, "loan_no"].ne(post_sold["loan_no"]).all(),
       "SF23: the receipts file shows %d Wirraway remittances on accounts sold in March 2026, after the sale; each belongs to "
       "the account then holding the reference" % len(post_sold))
    sold_true = asof[asof["loan_no"].isin(set(sale["loan_no"])) & (asof["received_date"] > cut.isoformat())]
    ok(len(sold_true) == 0, "nothing is received on a sold account after the cut-off once each remittance is given its placement")
    quiet = asof[asof["loan_no"].isin(set(sale["loan_no"])) & (asof["received_date"] > (cut - dt.timedelta(days=W.SALE_QUIET_DAYS)).isoformat())
                 & (asof["received_date"] <= cut.isoformat())]
    ok(len(quiet) == 0, "sold accounts had nothing received in the six months to the cut-off")
    x = asof.merge(fy[["loan_no", "default_month", "resolution", "resolution_date"]], on="loan_no", how="inner")
    me = pd.to_datetime(x["default_month"] + "-01") + pd.offsets.MonthEnd(0)
    ok((pd.to_datetime(x["received_date"]) > me).all(), "every receipt on a 2023-24 default falls after its default month end")
    ag2 = x[x["source"] == "Agency remittance"]
    ok(ag2["loan_no"].isin(set(fy.loc[fy["resolution"] == "Written off", "loan_no"]) | stale).all(),
       "agency remittances fall on written-off loans only, those closed as cured on a withdrawn arrangement included")
    wof = pd.read_csv(os.path.join(T, "write_offs_2025-07_2026-09.csv"))
    mar = float(wof.loc[wof["month"] == "2026-03", "recoveries_aud"].iloc[0])
    other = wof.loc[wof["month"] != "2026-03", "recoveries_aud"].astype(float)
    ok(mar - float(blank["amount_aud"].iloc[0]) > 0 and other.between(150000, 600000).all(),
       "the March 2026 recoveries carry the sale on top of an ordinary month")
    return dict(defaults=len(fy), receipts=len(led), sold=len(sale), consideration=float(blank["amount_aud"].iloc[0]),
                fy_sold=len(fs), reversals=len(neg), withdrawn_written_off=len(stale), withdrawn_caught_up=len(caught),
                cured_label=int((fy["resolution"] == "Cured").sum()), placements=len(pl), reissued=int((many > 1).sum()),
                misposted=int(moved.sum()), post_sale=len(post_sold))


# ---------------------------------------------------------------------------------------------------------- ask A

# the ask A ladder, each stop a complete path: (name, cures, sale, disc, refs). The careful path executes 3.3 as filed
# (the sale shared by face value, receipts discounted monthly from the default month end, dishonours netted) and takes
# the receipts file's loan numbers and the Collections cures as they stand
ASK_A_STOPS = [("careful: cures at their word, remittances as posted", "label", "share", "monthly", "posted"),
               ("A1 cracked, remittances as posted", "gold", "share", "monthly", "posted"),
               ("A2 cracked, cures at their word", "label", "share", "monthly", "asof"),
               ("over-cleaner: every withdrawn cure doubted, as posted", "no_withdrawn", "share", "monthly", "posted"),
               ("A2 to the reference's latest placement", "gold", "share", "monthly", "latest"),
               ("A2 reissued references dropped", "gold", "share", "monthly", "drop"),
               ("latest placement, cures at their word", "label", "share", "monthly", "latest"),
               ("the sale left out", "gold", "none", "monthly", "asof"),
               ("the sale spread over the written-off cases", "gold", "spread", "monthly", "asof"),
               ("not discounted", "gold", "share", "none", "asof"),
               ("the sale left out, cures at their word, as posted", "label", "none", "none", "posted"),
               ("every cure measured on its receipts", "receipts", "share", "monthly", "asof")]
ASK_A_MISREADINGS = [("days", "days"), ("month_start", "month_start"), ("annual", "annual")]


def ask_a():
    g = E.realised_lgd(T)
    ok(len(g) == 8, "eight figures, two products by four quarters")
    ok(all(40 <= v <= 95 for v in g.values()), "realised losses sit where unsecured personal lending loses")
    for key, v in g.items():
        ok(edge(v, 0.1) >= 0.05, "realised loss %s Q%d clears its rounding edge by 5 per cent of the bin (%.3f)" % (key[0], key[1], edge(v, 0.1)))
    ok(all(abs(E.lgd_path(T)[k] - g[k]) < 1e-9 for k in g), "the composable path reproduces the answer")
    vg = V.lgd(T)
    ok(all(abs(vg[k] - g[k]) < 1e-9 for k in g), "the independent verifier agrees on all eight")
    stops = {}
    for name, c, sl, dc, rf in ASK_A_STOPS:
        x = E.lgd_path(T, cures=c, sale=sl, disc=dc, refs=rf)
        stops[name] = x
        near = min(abs(x[k] - g[k]) for k in g)
        ok(all(round(x[k], 1) != round(g[k], 1) for k in g) and near >= 0.05,
           "ask A stop %s misses every figure at one decimal (nearest %.2f points)" % (name, near))
    for name, v in ASK_A_MISREADINGS:
        x = E.realised_lgd(T, v, "gold")
        stops[name] = x
        hit = sum(round(x[k], 1) == round(g[k], 1) for k in g)
        ok(hit <= 2, "SF21: 3.3 misread as %s misses %d of the 8 figures" % (name, 8 - hit))
    for name, kw in (("cures at their word", dict(cures="label")), ("remittances as posted", dict(refs="posted")),
                     ("both", dict(cures="label", refs="posted"))):
        x, vx = E.lgd_path(T, **kw), V.lgd(T, **kw)
        ok(all(abs(vx[k] - x[k]) < 1e-9 for k in g), "the verifier reproduces ask A stop %s" % name)
    for v in ("floor", "face_alloc"):
        x = E.realised_lgd(T, v, "gold")
        ok(all(round(x[k], 1) == round(g[k], 1) and abs(x[k] - g[k]) < 0.01 for k in g), "%s gives the same eight figures (converged)" % v)
    return g, stops


# ---------------------------------------------------------------------------------------------------------- ask B

ASK_B_STOPS = [("June days past due as extracted", dict(june_dpd="extract")),
               ("1 labels, May run", dict(june_rule="label", run="may", june_dpd="extract")),
               ("2 dates without any cancelled, May run", dict(june_rule="excl_cancel", run="may", june_dpd="extract")),
               ("3 dates, May run", dict(run="may", june_dpd="extract")),
               ("4 dates without any cancelled, June run", dict(june_rule="excl_cancel", june_dpd="extract")),
               ("5 cancellation ignored, June run", dict(june_rule="ignore_cancel", june_dpd="extract")),
               ("6 September scores for June", dict(run="sept", june_dpd="extract")),
               ("labels, June run", dict(june_rule="label", june_dpd="extract")),
               ("labels with start dates, May run", dict(june_rule="label_dated", run="may", june_dpd="extract")),
               ("July run", dict(run="july", june_dpd="extract")),
               ("end date excluded", dict(june_rule="end_exclusive", june_dpd="extract")),
               ("the 1 July run as it stands", dict(run="june1", june_dpd=1)),
               ("the 1 July run, Scored rows only", dict(run="june1_scored", june_dpd=1)),
               ("the 1 July run as it stands, June days past due as extracted", dict(run="june1", june_dpd="extract"))]


def ask_b():
    m = E.migration(T)
    ok(len(m) == 16 and all(v > 0 for v in m.values()), "sixteen counts, none empty")
    vm = V.migration(T)
    ok(vm == m, "the independent verifier agrees on all sixteen")
    hd = pd.read_excel(os.path.join(T, E.F["hardship"]), sheet_name=0)
    j = pd.Timestamp(JUNE)
    s_, e_, c_ = hd["arrangement_start"], hd["arrangement_end"], hd["cancelled_date"]
    inf = (s_ <= j) & (e_ >= j) & (c_.isna() | (c_ > j))
    facts = dict(in_force=int(inf.sum()), still_active=int((inf & (hd["arrangement_status"] == "Active")).sum()),
                 cancelled_after=int((inf & c_.notna()).sum()), end_on_30_june=int((inf & (e_ == j)).sum()),
                 cancelled_on_30_june=int((c_ == j).sum()))
    ok(facts["in_force"] >= 500 and facts["still_active"] <= 0.3 * facts["in_force"] and facts["cancelled_after"] >= 30
       and facts["end_on_30_june"] >= 1 and facts["cancelled_on_30_june"] == 0, "SF28: the register at 30 June: %s" % facts)
    rr = pd.Timestamp(R)
    old_rule = set(hd.loc[(hd["arrangement_status"] != "Cancelled") & (s_ <= rr) & (e_ >= rr), "loan_no"])
    ok(E.in_force(hd, R, "dates") == old_rule, "2.2 read with or without the cancelled rows gives one set at 30 September")
    stops = {}
    for name, kw in ASK_B_STOPS:
        x = E.migration(T, **kw)
        stops[name] = x
        moved = sum(x[k] != m[k] for k in m)
        ok(moved >= 6, "ask B stop %s lands off the answer (%d of 16 move)" % (name, moved))
    for name in ("June days past due as extracted", "1 labels, May run", "2 dates without any cancelled, May run",
                 "3 dates, May run", "6 September scores for June"):
        ok(sum(stops[name][k] != m[k] for k in m) >= 12, "stop %s moves at least 12 of the 16 counts" % name)
    names = [n for n, _ in ASK_B_STOPS]
    for a_ in range(len(names)):
        for b_ in range(a_ + 1, len(names)):
            ok(stops[names[a_]] != stops[names[b_]], "stops %r and %r land on different counts" % (names[a_], names[b_]))
    same = E.migration(T, june_rule="cancel_inclusive")
    ok(same == m, "reading a cancellation on 30 June either way gives the same counts (none falls on it)")
    hz = E.migration(T, sept_rule="label")
    hz_moved = sum(hz[k] != m[k] for k in m)
    ok(hz_moved <= 2, "staging September on the status label moves %d counts (the status lag, a hazard)" % hz_moved)
    for nm, kw, ref in (("June days past due as extracted", dict(dpd="extract"), "June days past due as extracted"),
                        ("labels, May run", dict(june="label", run="2026-06-01", dpd="extract"), "1 labels, May run"),
                        ("every cancelled out, June run", dict(june="excl", dpd="extract"), "4 dates without any cancelled, June run")):
        ok(V.migration(T, **kw) == stops[ref], "the verifier reproduces ask B stop %s" % nm)
    return m, stops, facts, hz_moved


# ---------------------------------------------------------------------------------------------------------- ask forks (round 2)

# the round 2 judge's literal builds, as its verdict describes them: every non-sale receipt on the loan the receipts file
# posts it to, the sale by face value at 3.85 cents on 27 March 2026, monthly discounting by whole months, every Collections
# cure at its word, pooled by balance; the June stages from the June extract's days past due as they stand
JUDGE_R2_A = {("PL36", 1): 60.5, ("PL36", 2): 55.6, ("PL36", 3): 65.2, ("PL36", 4): 66.8,
              ("PL60", 1): 67.7, ("PL60", 2): 69.1, ("PL60", 3): 67.8, ("PL60", 4): 74.7}
JUDGE_R2_B = [14450, 680, 39, 1582, 488, 658, 118, 124, 17706, 830, 43, 1176, 578, 943, 143, 113]
NOW = ("stage 1", "stage 2", "stage 3", "closed")
GROUPS = [("PL36", 1), ("PL36", 2), ("PL60", 1), ("PL60", 2)]


def ask_forks(s, g_a, m_b):
    """SF22, SF23, SF25 and SF29 to SF33: every reading the asks' wording and the filed text leave, closed by a quoted
    sentence or converging; and the judge's round 2 builds reproduced, each parting from the golden only where the repaired
    filed text now rules"""
    keys = sorted(g_a)
    hit_a = lambda x: sum(round(x[k], 1) == round(g_a[k], 1) for k in keys)
    flat = lambda m: [m[(p, s0, x)] for p, s0 in GROUPS for x in NOW]
    pol = docx_text(POLICY())
    out = {}
    # the judge's round 2 builds reproduce its verdict to the figure
    ja = E.lgd_path(T, cures="label", refs="posted")
    ok(all(round(ja[k], 1) == JUDGE_R2_A[k] for k in keys), "the judge's round 2 loss-share build reproduces its eight figures")
    jb_ = E.migration(T, june_dpd="extract")
    ok(flat(jb_) == JUDGE_R2_B, "the judge's round 2 slide 2 build reproduces its sixteen counts")
    # each step where it parts from the golden is now ruled on by a filed sentence, and changing only that step walks to it
    a1 = E.lgd_path(T, cures="gold", refs="posted")
    a2 = E.lgd_path(T, cures="label", refs="asof")
    ok(hit_a(a1) == 0 and hit_a(a2) == 0 and hit_a(E.lgd_path(T)) == 8, "SF22, SF23: the judge's build parts from the golden "
       "at two steps only, the cure and the attribution; either one alone still misses all eight, both together give all eight")
    ok(flat(E.migration(T)) == flat(m_b), "SF25: the judge's slide 2 build with only the June clock read from the change log "
                                          "gives all sixteen counts")
    # SF22: 3.3's cure sentence, executed on the loan record, selects the same cases the hardship register does (C1)
    fy, led, _, _, _ = E.collections(T)
    withdrawn, stale = E.withdrawn_cures(T, fy, led)
    gone = E.record_gone(T)
    cured_label = set(fy.loc[fy["resolution"] == "Cured", "loan_no"])
    ok(cured_label & gone == stale and len(stale) >= 30, "SF22: the cases Collections closed as cured that the record shows "
       "written off or settled for less (%d) are exactly those cured on an arrangement withdrawn before it started" % len(stale))
    rec = pd.read_parquet(os.path.join(T, E.F["record"]), columns=["loan_no", "closure_reason", "status"])
    rc = rec[rec["loan_no"].isin(cured_label - gone)]
    ok(set(rc["closure_reason"].fillna("open")) <= {"Repaid in full", "open"}, "SF22: every other cure is repaid in full or still "
       "open (%d), so 3.3's sentence moves no genuine cure" % len(rc))
    reg = E.realised_lgd(T, cures="register")
    ok(all(abs(reg[k] - g_a[k]) < 1e-9 for k in keys), "SF22: the cure read through the register gives the same eight figures (C1)")
    # SF23: what the receipts file's loan number and the agency's reference hold, and the placement report's dates
    pl = E.placements(T)
    ag = led[led["source"] == "Agency remittance"].copy()
    ag["ref"] = ag["narrative"].str.extract(r"ref (\d+)$")[0].astype(int)
    ag["day"] = pd.to_datetime(ag["received_date"])
    m = ag.reset_index().merge(pl[["Agency Ref", "Client Ref", "placed", "closed"]], left_on="ref", right_on="Agency Ref")
    win = m[(m["placed"] <= m["day"]) & (m["closed"].isna() | (m["day"] <= m["closed"]))]
    ok(win["index"].is_unique and len(win) == len(ag), "SF23: every one of the %d remittances falls inside exactly one placement "
       "of the reference it quotes" % len(ag))
    edge_days = int(((m["day"] == m["placed"]) | (m["day"] == m["closed"])).sum())
    asof = m[m["placed"] <= m["day"]].sort_values("placed").groupby("index").tail(1).set_index("index")["Client Ref"]
    ok(edge_days == 0 and (asof.sort_index() == win.set_index("index")["Client Ref"].sort_index()).all(),
       "SF23: no remittance falls on a placement's first or last day, and the placement last placed under the reference on "
       "or before the day gives the same account (C1)")
    own = ag.reset_index().merge(pl[["Agency Ref", "Client Ref", "placed", "closed"]], left_on=["ref", "loan_no"],
                                 right_on=["Agency Ref", "Client Ref"], how="left")
    outside = own[own["placed"].isna() | (own["day"] < own["placed"]) | (own["closed"].notna() & (own["day"] > own["closed"]))]
    moved = int((win.set_index("index")["Client Ref"] != ag["loan_no"]).sum())
    ok(len(outside) == moved >= 500, "SF23: each of the %d remittances posted to another account than the one placed under its "
       "reference that day arrives outside the posted account's own placement, so the posted reading has Wirraway remitting "
       "on accounts it no longer held" % moved)
    out["misposted"] = moved
    # SF25: the June clock. The dictionary counts the June extract's days past due to its run date, 1 July, and the
    # September extract's to the reporting date; 1.3 stages on close of business at the reporting date. Executed literally
    # (each June value put back by the days between the two dates the dictionary gives) the sixteen counts are the golden;
    # the column as it stands misses. The record and 1.4 bear the clock out
    cj, cs = V.dpd_clock(T, E.F["june"]), V.dpd_clock(T, E.F["book"])
    ok(cj == (JUNE, JUNE + dt.timedelta(days=1)) and cs == (dt.date(2026, 9, 30), dt.date(2026, 9, 30)),
       "SF25: the dictionary's lines give June's days past due counted to 1 July for loans open at close of business on 30 "
       "June, and September's counted to 30 September")
    jx = pd.read_csv(os.path.join(T, E.F["june"]))
    lit_m = E.migration(T, june_dpd=(cj[1] - cj[0]).days)
    ok(flat(lit_m) == flat(m_b), "SF25: the literal reading with the dictionary's June clock gives all sixteen golden counts")
    ok(sum(a != b for a, b in zip(flat(jb_), flat(m_b))) >= 12 and flat(jb_) == JUDGE_R2_B,
       "SF25: the column read unadjusted misses (the judge's literal row)")
    lr = pd.read_parquet(os.path.join(T, E.F["record"]), columns=["loan_no", "default_month"])
    x90 = jx[jx["days_past_due"] == 90].merge(lr, on="loan_no")
    ok(len(x90) >= 3 and (x90["default_month"] == "2026-07").all(), "SF25: the %d loans the June extract shows at 90 days were at "
       "89 at close of business, in stage 2 at 30 June, and the record dates each default to the July month end" % len(x90))
    ok(int((jx["days_past_due"] >= 180).sum()) == int((jx["days_past_due"] == 180).sum()) >= 5, "SF25: the loans the June "
       "extract shows at 180 days were at 179 when the June month-end batch ran, so 1.4 had not written them off")
    moved_b = sum(a != b for a, b in zip(flat(jb_), flat(m_b)))
    ok(moved_b >= 12, "SF25: the June column read unadjusted moves %d of 16" % moved_b)
    # SF34: the June behaviour score. 2.1 compares the score at the reporting date; the scoring file dates each run by the
    # day it was made and marks a score the run could not calculate as Carried forward (the last score held). Read
    # literally, the June month end's score is its Scored row: the 1 July run's, or the 2 July rerun's for the partition the
    # 1 July run carried forward. Each loan's latest June month-end row gives the same (C1); the run dated 1 July taken as
    # it stands, or its Scored rows alone, miss
    meta_s = meta_of(os.path.join(T, E.F["scores"]))
    ok(meta_s["score_status"].startswith("Scored, or Carried forward (the run could not score the loan and repeats the last "
                                         "score the scoring engine holds for it)"), "SF34: the scoring file says what a "
                                                                                    "Carried forward score is")
    ok(E.migration(T, run="june_latest") == m_b, "SF34: each loan's latest June month-end row gives all sixteen counts (C1)")
    for nm_, run_ in (("the run dated 1 July as it stands", "june1"), ("its Scored rows only", "june1_scored")):
        x = E.migration(T, run=run_)
        mv = sum(x[k] != m_b[k] for k in m_b)
        ok(mv >= 12, "SF34: %s moves %d of 16, ruled out by the file's own status and run dates" % (nm_, mv))
        out["b4 " + nm_] = mv
    vx = V.migration(T, run="2026-07-01")
    ok(sum(vx[k] != m_b[k] for k in m_b) >= 12, "SF34: the verifier's own build on the run dated 1 July as it stands misses")
    # SF31: 2.2, cancelled arrangements and the end date, at 30 June
    dic = " ".join(open(os.path.join(T, "ecl_extract_dictionary.txt")).read().split())
    stops_day = "one that is cancelled stops being in force on the day it is cancelled"
    for name, kw, quote, where in (("every cancelled arrangement out, whenever cancelled", dict(june_rule="excl_cancel"), stops_day, pol),
                                   ("the cancellation ignored", dict(june_rule="ignore_cancel"), stops_day, pol),
                                   ("the end date excluded", dict(june_rule="end_exclusive"),
                                    "last day of the varied repayments, as approved", dic)):
        x = E.migration(T, **kw)
        mv = sum(x[k] != m_b[k] for k in m_b)
        ok(mv >= 6 and quote in where, "SF31: %s moves %d of 16, ruled out by %r" % (name, mv, quote))
        out[name] = mv
    ok(E.migration(T, june_rule="cancel_inclusive") == m_b, "SF31: no cancellation falls on 30 June, so the cancellation day "
                                                            "reads either way (C1)")
    # SF29: pooled by balance against the mean of the loans' own shares
    mean = E.realised_lgd(T, pool="mean")
    ok(hit_a(mean) == 0, "SF29: the mean of each loan's own share misses all eight (nearest %.2f), ruled out by 3.3's share summed "
       "over the quarter's defaults" % min(abs(mean[k] - g_a[k]) for k in keys))
    out["mean_nearest"] = min(abs(mean[k] - g_a[k]) for k in keys)
    # SF30: the realised losses in place of the filed LGDs would move the allowance; 3.3 keeps 70 and 76 in force
    fyq = fy.assign(B=fy["balance_at_default"].astype(float), q=[((int(mm[5:7]) - 7) % 12) // 3 + 1 for mm in fy["default_month"]])
    bal = fyq.groupby(["product", "q"])["B"].sum()
    by_p = {p: sum(g_a[(p, q)] * bal[(p, q)] for q in (1, 2, 3, 4)) / sum(bal[(p, q)] for q in (1, 2, 3, 4)) / 100 for p in ("PL36", "PL60")}
    L = s["loans"]
    alt = float(sum(L.loc[L["product"] == p, "ecl"].sum() * by_p[p] / E.LGD[p] for p in by_p))
    ok(M(alt) != M(s["total"]) and "the rates in force until the Audit and Risk Committee approves a change" in pol,
       "SF30: the realised losses (PL36 %.1f, PL60 %.1f per cent) in place of 70 and 76 would book AUD %.1f million, ruled out by "
       "3.3's rates in force" % (100 * by_p["PL36"], 100 * by_p["PL60"], M(alt)))
    out["lgd_swap_allowance"] = alt
    # SF32: the 60-point fall on integer scores, the converted v3 score rounded as the migration note states
    note = " ".join(p_.extract_text() for p_ in E.PdfReader(os.path.join(T, E.F["note"])).pages)
    ok("v4 = 1.12 x v3 - 18, rounded to the nearest whole point" in " ".join(note.split()), "SF32: the migration note rounds the "
                                                                                           "converted score to a whole point")
    frac = ((1.12 * jx.loc[jx["origination_scorecard"] == "v3", "origination_score"] - 18) % 1 - 0.5).abs()
    sc = E.score_runs(T)
    ok(frac.min() >= 0.015 and (sc["score_v4"] % 1 == 0).all(), "SF32: no v3 score converts within %.2f of a half point and every "
       "score is a whole number, so each fall is a whole number of points" % frac.min())
    _, _, mapping = E.load_staging(T)
    a_, b_ = mapping
    o4 = np.where(jx["origination_scorecard"] == "v3", np.floor(a_ * jx["origination_score"] + b_ + 0.5), jx["origination_score"])
    jm_ = sc[(sc["run_date"] > JUNE) & (sc["run_date"] < dt.date(2026, 8, 1)) & (sc["score_status"] == "Scored")]
    js = jx["loan_no"].map(jm_.set_index("loan_no")["score_v4"])
    at60 = int(((o4 - js) == 60).sum())
    st_ = E.migration(T, strict_fall=True)
    mv = sum(st_[k] != m_b[k] for k in m_b)
    ok(at60 >= 30 and mv >= 6 and "60 or more points below" in pol, "SF32: %d June loans fall exactly 60 points and are in stage 2 "
       "under 2.1's 60 or more; reading it as more than 60 moves %d of 16" % (at60, mv))
    out["june_at_60"], out["strict_fall_moved"] = at60, mv
    # SF33: a written-off loan is closed, in no stage; read as stage 3 instead, the counts move
    bk = pd.read_csv(os.path.join(T, E.F["book"]))
    jc = jx.assign(days_past_due=(jx["days_past_due"] - 1).clip(lower=0))
    _, hd_, mp_ = E.load_staging(T)
    jst = E.stage(jc, hd_, mp_, at=JUNE, score=js)
    gone_q = jc.assign(st=jst.values)[~jc["loan_no"].isin(set(bk["loan_no"]))].merge(rec, on="loan_no")
    ok((gone_q["status"] == "Closed").all(), "SF33: every June loan missing from September is Closed in the record")
    grp = gone_q[gone_q["st"] < 3]
    wo_g = grp[grp["closure_reason"] == "Written off"].groupby(["product", "st"]).size()
    alt = dict(m_b)
    for (p_, s0), n_ in wo_g.items():
        alt[(p_, int(s0), "stage 3")] += int(n_)
        alt[(p_, int(s0), "closed")] -= int(n_)
    mv33 = sum(alt[k] != m_b[k] for k in m_b)
    ok(int(wo_g.sum()) >= 2 and mv33 >= 2, "SF33: %d loans of the four groups were written off by 30 September; counted in stage "
       "3 rather than closed they move %d of 16, ruled out by the dictionary's 'a closed loan is in no stage'" % (int(wo_g.sum()), mv33))
    out["closed_written_off"], out["written_off_as_stage3_moved"] = int(wo_g.sum()), mv33
    # the literal build of the repaired text, on the verifier's own code: the golden
    vl, vm = V.lgd(T), V.migration(T)
    ok(all(abs(vl[k] - g_a[k]) < 1e-9 for k in keys) and vm == m_b, "the literal execution of the filed text on the verifier's "
       "own code path gives all eight loss shares and all sixteen counts")
    return out


# ---------------------------------------------------------------------------------------------------------- separation

def separation(s):
    """the main call's rows: the September book's loans, the arrangements in force at 30 September, the loan record. No
    ask device and no hazard row sits inside them"""
    hd = pd.read_excel(os.path.join(T, E.F["hardship"]), sheet_name=0)
    j, r = pd.Timestamp(JUNE), pd.Timestamp(R)
    s_, e_, c_ = hd["arrangement_start"], hd["arrangement_end"], hd["cancelled_date"]
    june_only = (s_ <= j) & (e_ >= j) & (c_.isna() | (c_ > j)) & ~((s_ <= r) & (e_ >= r) & (c_.isna() | (c_ > r)))
    rows = hd[june_only]
    lab = rows[rows["arrangement_status"] == "Active"]
    canc = rows[rows["cancelled_date"].notna()]
    ok(len(lab) <= 2, "the arrangements in force at 30 June and not at 30 September (%d, %d cancelled after June) are labelled "
       "Active for %d (the status lag)" % (len(rows), len(canc), len(lab)))
    fy, led, _, _, _ = E.collections(T)
    withdrawn, stale = E.withdrawn_cures(T, fy, led)
    wd = hd[hd["loan_no"].isin(withdrawn) & (hd["arrangement_status"] == "Cancelled") & (hd["cancelled_date"] < hd["arrangement_start"])
            & (hd["received_date"] >= pd.Timestamp(2023, 7, 1))]
    bk = pd.read_csv(os.path.join(T, E.F["book"]))
    jb = pd.read_csv(os.path.join(T, E.F["june"]))
    covid = hd[(hd["hardship_reason"] == "COVID-19") & (hd["arrangement_status"] == "Active")]
    in_force = lambda d, x: x[(x["arrangement_start"] <= d) & (x["arrangement_end"] >= d) & (x["cancelled_date"].isna() | (x["cancelled_date"] > d))]
    on_book = set(s["loans"].loc[s["loans"]["stage"] < 3, "loan_no"])
    ok(len(in_force(r, wd)) == 0 and len(in_force(j, wd)) == 0 and len(in_force(r, covid)) == 0 and len(in_force(j, covid)) == 0,
       "no device row is in force at 30 September or 30 June: %d withdrawn collections arrangements, %d bulk COVID-19 rows"
       % (len(wd), len(covid)))
    ok(not covid["loan_no"].isin(set(bk["loan_no"]) | set(jb["loan_no"])).any() and not set(stale) & on_book,
       "no COVID-19 row and no written-off withdrawn cure sits on a loan in either book; the stage 1 and 2 loans hold none")
    counts = dict(june_only_rows=len(rows), cancelled_after_june=len(canc), withdrawn_arrangement_rows=len(wd),
                  stale_cure_cases=len(stale), covid_bulk_rows=len(covid), june_dpd_rows=int((jb["days_past_due"] > 0).sum()),
                  device_rows_in_main_population=0)
    tmp = tempfile.mkdtemp()
    try:
        cut = os.path.join(tmp, "main")
        shutil.copytree(T, cut)
        for kf in ASK_A_FILES + ASK_B_FILES:
            os.remove(os.path.join(cut, E.F[kf]))
        x = E.solve(cut)
        ok(all(x[n] == s[n] for n in FIG), "the main call never reads the June extract, the scoring runs, the Collections files, "
                                           "the sale schedule or the placement report")
        cut3 = os.path.join(tmp, "distractors")
        shutil.copytree(T, cut3)
        for f in DISTRACTORS:
            os.remove(os.path.join(cut3, f))
        x3 = E.solve(cut3)
        ok(x3["total"] == s["total"] and E.realised_lgd(cut3) == E.realised_lgd(T) and E.migration(cut3) == E.migration(T),
           "the two distractors are read by no graded figure")
    finally:
        shutil.rmtree(tmp)
    # the asks do not ride on the call: with the PD machinery disabled both asks still compute their figures
    keep = (E.hazards, E.forward_pd, E.landmark)

    def boom(*a, **k):
        raise RuntimeError("the PD was read")
    E.hazards = E.forward_pd = E.landmark = boom
    try:
        ga, mb = E.realised_lgd(T), E.migration(T)
    finally:
        E.hazards, E.forward_pd, E.landmark = keep
    ok(len(ga) == 8 and len(mb) == 16, "neither ask reads the PD, so landing or missing the call moves no ask figure")
    return counts


# ---------------------------------------------------------------------------------------------------------- readings and pins

STAGE2_HASH = "a8137550bc1822e3a354a8e4c71bfd98045823aa673609028826904547fc6d93"


def readings(s):
    """the golden's stage 2 set; the register's status lag; what the pause pin is worth"""
    L = s["loans"]
    ids = sorted(L.loc[L["stage"] == 2, "loan_no"])
    digest = hashlib.sha256("\n".join(ids).encode()).hexdigest()
    ok(len(ids) == 3266 and digest == STAGE2_HASH, "the stage 2 set is the one CRM-ECL-04 section 2 gives (%d loans, %s)" % (len(ids), digest[:12]))
    bk, hd, scen, mapping = E.load(T)
    end = hd["arrangement_end"].map(E._d)
    covid = (hd["hardship_reason"] == "COVID-19").to_numpy()
    stale = (hd["arrangement_status"] == "Active") & (end < R)
    ok(1 <= (stale & ~covid).sum() <= 5, "a few arrangements that ended after the last weekly status run still show Active (%d)"
       % (stale & ~covid).sum())
    last_run = R - dt.timedelta(days=(R.weekday() - W.STATUS_RUN_WEEKDAY) % 7)
    ok((end[stale & ~covid] >= last_run).all(), "each ended on or after the last Tuesday run, %s" % last_run)
    jb = pd.read_csv(os.path.join(T, E.F["june"]))
    bulk = stale & covid
    ok(bulk.sum() >= 300 and not hd.loc[bulk, "loan_no"].isin(set(bk["loan_no"]) | set(jb["loan_no"])).any(),
       "the bulk-loaded COVID-19 pauses of 2020 and 2021 still read Active (%d), every one on a loan closed before 30 June 2026" % bulk.sum())
    start = hd["arrangement_start"].map(E._d)
    canc = hd["arrangement_status"] == "Cancelled"
    want = [("Cancelled" if c else "Active" if cv else "Approved" if a > R else "Active" if b >= last_run else "Completed")
            for a, b, c, cv in zip(start, end, canc, covid)]
    ok((hd["arrangement_status"] == want).all(), "every other status agrees with the dates and the weekly run as at 30 September")
    pz = hd[(hd["arrangement_type"] == "Payment pause") & ~canc].assign(s=start, e=end)
    book = bk.set_index("loan_no")
    held = {}
    for loan, g in pz.groupby("loan_no"):
        if loan in book.index:
            f = book.at[loan, "first_rep"]
            n = sum(1 for d in due_dates(f, E.months_due(f)) if any(a <= d <= min(b, R) for a, b in zip(g["s"], g["e"])))
            if n:
                held[loan] = n
    md, le, cur = E.months_due, E.loan_ecl, {"n": 0}
    E.months_due = lambda f, report=E.REPORT: md(f, report) - cur["n"]

    def held_out(row, *a, **k):
        cur["n"] = held.get(row["loan_no"], 0)
        try:
            return le(row, *a, **k)
        finally:
            cur["n"] = 0
    E.loan_ecl = held_out
    try:
        alt = E.solve(T)
    finally:
        E.months_due, E.loan_ecl = md, le
    moved = [n for n in FIG if M(alt[n]) != M(s[n])]
    ok(len(held) > 1000 and not moved, "SF10: leaving pause dates out of the repayments fallen due (%d loans, %d dates) keeps all "
       "three figures in their bins (allowance %.0f): same call" % (len(held), sum(held.values()), alt["total"]))
    return (len(held), sum(held.values()), len(moved), alt["total"]), (int((stale & ~covid).sum()), int(bulk.sum()))


def residue_and_pins():
    bk = pd.read_csv(os.path.join(T, E.F["book"]))
    nulls = bk.isna().sum()
    ok(set(nulls[nulls > 0].index) == {"state"} and 20 <= nulls["state"] <= 100,
       "the book's only blanks are the state of borrowers with an overseas address (%d)" % nulls["state"])
    jb = pd.read_csv(os.path.join(T, E.F["june"]))
    jn = jb.isna().sum()
    ok(set(jn[jn > 0].index) == {"state"}, "the June extract's only blanks are overseas addresses too")
    hd = pd.read_excel(os.path.join(T, E.F["hardship"]), sheet_name=0)
    nr = hd[hd["hardship_reason"].isna()]
    early = hd[hd["received_date"].dt.date < W.REASON_REQUIRED]
    ok(3 <= len(nr) <= 0.25 * len(early) and (nr["received_date"].dt.date < W.REASON_REQUIRED).all(),
       "some cases from before the reason became a required field carry none (%d of %d)" % (len(nr), len(early)))
    lk = hd[(hd["hardship_reason"] == "COVID-19") & (hd["received_date"].dt.year == 2021)]
    ok(len(lk) > 50 and not lk["loan_no"].isin(set(bk["loan_no"])).any() and (lk["arrangement_end"] < "2022-03-01").all(),
       "the 2021 lockdown pauses (%d) sit on loans closed before the reporting date" % len(lk))
    pol = docx_text(POLICY())
    dic = " ".join(open(os.path.join(T, "ecl_extract_dictionary.txt")).read().split())
    ok("which stay in the loan's schedule" in dic and "the maturity date does not change" in dic
       and "a weekly batch on Tuesday nights moves an arrangement to Completed once its end date has passed" in dic,
       "payment pauses and the weekly status batch are described in the dictionary")
    ok("more than 30 days past due" in pol and "30 or more days past due" in pol, "2.5 quotes AASB 9 5.5.11 as more than 30 days")
    ok("last day of the varied repayments, as approved" in dic, "the dictionary dates an arrangement's end as the last day of the "
                                                                "varied repayments")
    ok("its behaviour score is 60 or more points below its score at origination" in pol and "scoring run on the position" not in pol,
       "2.1 reads the behaviour score at the reporting date")
    ok("Days past due are taken from the loan system at close of business on the reporting date." in pol
       and "once 180 or more days past due" in pol, "SF25: 1.3 takes days past due at close of business; 1.4 writes off at 180 in "
                                                    "the month-end batch")
    ok("month of the first month end at which the loan was 90 or more days past due" in dic
       and "month of the first month end at which the loan system recorded the loan 90 or more days past due" in dic,
       "the default month is a field definition, in the Collections file and the loan record")
    ok("including the proceeds of any sale of the debt" in pol and "count at the end of the month in which they are received" in pol
       and "discounted to the default month end monthly at the contract rate, as in 3.5" in pol
       and "A loan that returns to performing, its arrears made good or a collections hardship arrangement started, has lost "
           "nothing; a loan the loan record shows written off or settled for less did not return to performing and is "
           "measured on what it recovered." in pol,
       "SF21, SF22: 3.3 files the realised loss measure, when a loan has lost nothing and when a cure does not count")
    ok("for each product and quarter of default, the balance at default not recovered since, summed over the quarter's "
       "defaults, as a share of their total balance at default" in pol, "SF29: 3.3 files the share pooled by balance over each "
                                                                       "product and quarter of default")
    ok("LGD is 70 per cent for PL36 and 76 per cent for PL60, the rates in force until the Audit and Risk Committee approves "
       "a change." in pol, "SF30: 3.3 files 70 and 76 per cent as the rates in force for the allowance")
    ok("Version 4.2" in pol and "Effective 1 September 2026" in pol and "Annual review: 2.2 and 3.3 revised" in pol
       and "PD read each quarter from the loan record by month on book" in pol, "the version history files the record (3.0) and "
                                                                               "the LGD review (4.2)")
    ok("CHG-26-0418 21 Jul 2026 current_score_v4 added" in dic and dic.count("CHG-") == 6
       and "CHG-23-0207 6 Feb 2023 loan record added to the quarter-end run" in dic,
       "the dictionary keeps the job's change log, six entries, the record added in February 2023")
    ok("A loan leaves the extract once it closes: repaid in full, settled or written off; a closed loan is in no stage." in dic,
       "SF27, SF33: the dictionary says what closes a loan and that a closed loan, one written off included, is in no stage")
    ok("loan number the receipt was posted to; an agency remittance is posted on the Wirraway reference its narrative quotes" in dic
       and "Agency Ref Wirraway's reference for the account while it is placed" in dic,
       "SF23: the dictionary says what the receipts file's loan number and the agency's reference hold")
    ok("personal_loan_book_2026-06-30.csv run 1 July 2026: every loan open at close of business on 30 June 2026; "
       "days_past_due counted to the run date, 1 July 2026" in dic
       and "personal_loan_book_2026-09-30.csv run 1 October 2026: every loan open at close of business on 30 September "
           "2026; days_past_due counted to the reporting date, 30 September 2026" in dic
       and dic.count("days_past_due counted to") == 2,
       "SF25: the dictionary files each loan book extract's run clock once: June's days past due counted to the run date, "
       "September's to the reporting date")
    ok("current_score_v4 added; extract taken from the quarter-end end-of-day snapshot" in dic,
       "the job's change log records the extract's source since CHG-26-0418")
    for w in ("reissu", "first placement", "first placed", "loader", "re-used", "reused", "recycl", "wrong loan", "earlier holder"):
        ok(w not in dic.lower() and w not in pol.lower(), "the filed text says what each field holds and narrates no reissue or "
                                                          "posting rule (%r)" % w)
    for w in ("rerun", "re-run", "run again", "partition", "carried forward", "2 july"):
        ok(w not in dic.lower() and w not in pol.lower(), "the dictionary and the policy say nothing of the June month end's "
                                                          "scoring rerun (%r)" % w)
    for w in ("a day more", "one day more", "day further", "clock", "subtract", "put back", "minus"):
        ok(w not in dic.lower() and w not in pol.lower(), "the dictionary records each extract's run clock and says nothing "
                                                          "of recounting days past due (%r)" % w)
    for w in ("whatever resolution", "already reflects a loan's recent repayment conduct", "confirms it at each year end"):
        ok(w not in pol.lower(), "CRM-ECL-04 states 2.4, 3.3 and 3.6 without pre-emptive justification (%r)" % w)
    for f in sorted(os.listdir(T)):
        if not f.endswith((".docx", ".xlsx")):
            continue
        with zipfile.ZipFile(os.path.join(T, f)) as z:
            app = "word" if f.endswith(".docx") else "excel"
            th = z.read("%s/theme/theme1.xml" % ("word" if app == "word" else "xl")).decode("utf-8")
            core = z.read("docProps/core.xml").decode()
            ok(th == OX.theme_xml(app), "%s carries the Office 2013-2022 theme as %s writes it" % (f, app.title()))
            made = re.search(r"<dcterms:created[^>]*>(\d{4}-\d\d-\d\d)", core).group(1)
            who = re.search(r"<dc:creator>([^<]*)", core).group(1)
            if app == "excel":
                ok('defaultThemeVersion="166925"' in z.read("xl/workbook.xml").decode(), "%s: workbook theme version 166925" % f)
                ok(made < "2023-07-01" or who == W.JOB_ACCOUNT, "%s: created %s, before the 2023 theme, or by the job" % (f, made))
            if f == E.F["hardship"]:
                ok(who == W.JOB_ACCOUNT and "<cp:lastModifiedBy>Declan Wallace<" in core, "the register is written by the job's account")
    with zipfile.ZipFile(POLICY()) as z:
        app = z.read("docProps/app.xml").decode()
        doc = z.read("word/document.xml").decode()
    st = {k: int(v) for k, v in re.findall(r"<(Lines|Paragraphs|Words|Characters|CharactersWithSpaces)>(\d+)<", app)}
    ok(st["Lines"] >= 1.25 * st["Paragraphs"] and st["CharactersWithSpaces"] > st["Characters"] > 4 * st["Words"],
       "the policy's statistics are the word processor's measurement (%s)" % st)
    rs = Counter(re.findall(r'<w:p [^>]*w:rsidR="([0-9A-F]{8})"', doc))
    ok(len(rs) >= 8 and rs.most_common(1)[0][1] >= 0.25 * sum(rs.values()) and doc.count('<w:r w:rsidR="') >= 10,
       "the policy's revision ids follow its history (%d sessions)" % len(rs))
    ok("Version 4.2" in pol and "25 August 2026" in pol, "the policy is version 4.2, approved in August 2026")
    pl = pd.read_csv(os.path.join(T, E.F["placements"]), dtype=str)
    ok(pl["Placed Balance"].astype(float).min() > 0 and pl["Closed"].isna().sum() > 500, "placements carry their balances; "
       "%d still open" % pl["Closed"].isna().sum())
    raw = open(os.path.join(T, "ecl_extract_dictionary.txt")).read()
    named = [f for f in sorted(os.listdir(T)) if f != "ecl_extract_dictionary.txt" and f not in raw]
    ok(not named, "H9: every shipped file is entered in the run folder's dictionary by name (missing: %s)" % named)
    meta = json.load(open(os.path.join(TASK, "metadata.json")))
    ok(sorted(x["path"] for x in meta["input_files"]) == sorted(os.listdir(T)), "H9: metadata lists exactly the shipped files")


# ---------------------------------------------------------------------------------------------------------- ask-layer audits

REC_W, IF_W, ASK_W, R_KEEP = 38.0, 7.0, 55.0, 5.0


def audits(s, out, g_a, m_b):
    res = {}
    keys = sorted(g_a)
    hit_a = lambda x: sum(round(x[k], 1) == round(g_a[k], 1) for k in keys)
    hit_b = lambda x: sum(x[k] == m_b[k] for k in m_b)
    lazy_a = E.realised_lgd(T, "lazy", "label", refs="posted")
    lazy_b = E.migration(T, june_rule="label", run="may", june_dpd="extract")
    ok(hit_a(lazy_a) == 0, "the lazy path misses every ask A figure")
    ok(16 - hit_b(lazy_b) >= 12, "the lazy path misses at least 12 ask B counts")
    res["lazy"] = (8 - hit_a(lazy_a), 16 - hit_b(lazy_b))
    mat = {}
    for name, kw in (("A1 (filed in 3.3): cases closed as cured taken at their word", dict(cures="label")),
                     ("A2 (silent): remittances on the loan the receipts file posts them to", dict(refs="posted")),
                     ("A: the sale left out", dict(sale="none")), ("A: receipts not discounted", dict(disc="none")),
                     ("A: every cure measured on its receipts", dict(cures="receipts")),
                     ("A: the sale receipt spread over every write-off", dict(sale="spread"))):
        x = E.lgd_path(T, **kw)
        mat[name] = 8 - hit_a(x)
        ok(mat[name] == 8, "%s moves all eight ask A figures" % name)
    for name, kw in (("B4: the run dated 1 July as it stands", dict(run="june1")),
                     ("B4: the run dated 1 July, Scored rows only", dict(run="june1_scored")),
                     ("B3 (filed): June days past due as extracted", dict(june_dpd="extract")),
                     ("B1: the label after the switch", dict(june_rule="label")),
                     ("B1: every cancelled out", dict(june_rule="excl_cancel")),
                     ("B2: the run made before the date", dict(run="may"))):
        x = E.migration(T, **kw)
        mat[name] = 16 - hit_b(x)
        ok(mat[name] >= 12, "%s moves %d of the 16 ask B counts" % (name, mat[name]))
    res["matrix"] = mat
    worst = {1: 0, 2: 0, 3: 0, 4: 0}
    where = {}
    gold_path = ("gold", "share", "monthly", "asof")
    for c in ("gold", "label", "no_withdrawn", "receipts"):
        for sl in ("share", "none", "spread", "flat"):
            for dc in ("monthly", "none", "annual", "days", "month_start"):
                for rf in ("asof", "posted", "latest", "drop"):
                    path = (c, sl, dc, rf)
                    if path == gold_path:
                        continue
                    n_wrong = sum(x != y for x, y in zip(path, gold_path))
                    h = hit_a(E.lgd_path(T, c, sl, dc, rf))
                    if dc == "days" and n_wrong == 1:
                        continue
                    if h > worst[n_wrong]:
                        worst[n_wrong], where[n_wrong] = h, path
    ok(worst[1] <= 1 and worst[2] <= 1 and worst[3] <= 2, "composed ask A paths: one or two mishandlings reach at most %d of 8, "
       "three at most %d (%s)" % (max(worst[1], worst[2]), worst[3], where))
    res["composed_a"] = worst
    res["composed_a_where"] = where
    worst_b = 16
    for dpd in ("close", "extract"):
        for rule in ("dates", "label", "excl_cancel", "ignore_cancel", "end_exclusive"):
            for run in ("june", "june1", "june1_scored", "may", "july", "sept"):
                if (dpd, rule, run) == ("close", "dates", "june"):
                    continue
                worst_b = min(worst_b, 16 - hit_b(E.migration(T, june_rule=rule, run=run, june_dpd=dpd)))
    ok(worst_b >= 6, "composed ask B paths: every one moves at least %d of the 16 counts" % worst_b)
    res["composed_b_min_moved"] = worst_b
    # the pair simulation. The cracker lands the call, runs the hygiene battery and executes every filed rule it points at
    # (3.3's measure: the share pooled by balance, the sale by face value, receipts discounted monthly from the default
    # month end, a cure the record shows written off measured on its recoveries; dishonours netted; hardship by 2.2's
    # dates; the June days past due put back to 30 June by the dictionary's run clock). The battery's joins all land: the
    # receipts to the defaults, the record and the sale schedule, so it takes the remittances as posted (A2); the June
    # extract to the scoring run dated 1 July, every loan past its first repayment scored, so it takes that run as it
    # stands, carried-forward scores included (B4). The mirror stops on the default path and takes the same ask paths
    mirror_main = out[MAIN[0][0]]
    ok(all(M(mirror_main[n]) != M(s[n]) for n in FIG), "the mirror sheet misses the call")
    careful_a = E.lgd_path(T, cures="gold", refs="posted")
    careful_b = E.migration(T, june_dpd=1, run="june1")
    crack_b4 = E.migration(T, june_dpd=1)
    c_hits = hit_a(careful_a) + hit_b(careful_b)
    m_hits = c_hits
    ok(hit_a(careful_a) == 0 and hit_b(careful_b) <= 1 and hit_b(crack_b4) == 16, "pair simulation: the careful path keeps "
       "%d of 8 loss shares and %d of 16 counts; with the June month end's rerun read as well, all 16" % (hit_a(careful_a), hit_b(careful_b)))
    cracker = REC_W + IF_W + ASK_W * c_hits / 24.0
    mirror = R_KEEP + IF_W + ASK_W * m_hits / 24.0
    pair = (cracker + mirror) / 2
    ok(pair <= 40.0, "pair simulation: cracker %.1f (%d of 24), mirror %.1f (%d of 24), pair %.1f" % (cracker, c_hits, mirror, m_hits, pair))
    ok(100.0 * m_hits / 24 < 35.0, "mirror leakage %.1f per cent of the ask weight" % (100.0 * m_hits / 24))
    both_on = (2 * (REC_W + IF_W) + ASK_W * (c_hits + m_hits) / 24.0) / 2
    ok(both_on < 50.0, "both top responses on the call, each executing the June clock: pair %.1f" % both_on)
    # a mirror on rung 1 or 2 lands stage 2's bin by offsetting errors and may keep that component too
    mirror_s2 = R_KEEP + 5.0 + IF_W + ASK_W * m_hits / 24.0
    ok((cracker + mirror_s2) / 2 <= 40.0, "pair with a mirror on rung 1 or 2 keeping stage 2: %.1f" % ((cracker + mirror_s2) / 2))
    # the most obvious device is A2 (the dictionary's loan_no entry names the posting); the cracker cracking it
    a2_crack = REC_W + IF_W + ASK_W * (8 + hit_b(careful_b)) / 24.0
    ok((a2_crack + mirror) / 2 <= 40.0, "the cracker cracking A2, the most obvious device: pair %.1f" % ((a2_crack + mirror) / 2))
    b4_crack = REC_W + IF_W + ASK_W * (hit_a(careful_a) + 16) / 24.0
    print("RECORD residuals: the cracker cracking B4 instead %.1f, both A2 and B4 %.1f; both top responses cracking A2 %.1f, "
          "B4 %.1f" % ((b4_crack + mirror) / 2, (REC_W + IF_W + ASK_W + mirror) / 2,
                        (2 * (REC_W + IF_W) + ASK_W * 2 * (8 + hit_b(careful_b)) / 24.0) / 2,
                        (2 * (REC_W + IF_W) + ASK_W * 2 * (hit_a(careful_a) + 16) / 24.0) / 2))
    res["pair"] = dict(cracker=cracker, mirror=mirror, pair=pair, cracker_hits=c_hits, mirror_hits=m_hits,
                       mirror_leakage=100.0 * m_hits / 24, both_on_call=both_on,
                       cracker_cracks_a2=(a2_crack + mirror) / 2, cracker_cracks_b4=(b4_crack + mirror) / 2,
                       cracker_cracks_both=(REC_W + IF_W + ASK_W + mirror) / 2,
                       both_crack_a2=(2 * (REC_W + IF_W) + ASK_W * 2 * (8 + hit_b(careful_b)) / 24.0) / 2,
                       both_crack_b4=(2 * (REC_W + IF_W) + ASK_W * 2 * (hit_a(careful_a) + 16) / 24.0) / 2,
                       mirror_keeping_stage2=(cracker + mirror_s2) / 2)
    habit_a = E.lgd_path(T, cures="label", sale="none", disc="none", refs="posted")
    habit_b = E.migration(T, june_rule="excl_cancel", run="may", june_dpd="extract")
    h_hits = hit_a(habit_a) + hit_b(habit_b)
    h_pair = (REC_W + IF_W + R_KEEP + IF_W + 2 * ASK_W * h_hits / 24.0) / 2
    ok(h_pair <= 40.0, "pair simulation, habitual solvers: pair %.1f (%d of 24)" % (h_pair, h_hits))
    res["pair"]["habitual_pair"] = h_pair
    oc = {"A: every cure on a withdrawn arrangement doubted": E.lgd_path(T, cures="no_withdrawn", refs="posted"),
          "A: remittances on reissued references dropped": E.lgd_path(T, refs="drop"),
          "A: the sale receipt spread over every write-off": E.lgd_path(T, sale="spread")}
    for name, x in oc.items():
        ok(hit_a(x) == 0, "%s misses every ask A figure" % name)
    nr = E.realised_lgd(T, "no_reversals", "gold")
    res["no_reversals_moved"] = 8 - hit_a(nr)
    ok(res["no_reversals_moved"] == 0, "the dishonour reversals are texture: dropping them as errors moves no ask A figure")
    import openpyxl
    texts = docx_text(POLICY()) + open(os.path.join(T, "ecl_extract_dictionary.txt")).read()
    for f in sorted(os.listdir(T)):
        if f.endswith(".xlsx"):
            wb = openpyxl.load_workbook(os.path.join(T, f), read_only=True)
            for ws in wb.worksheets:
                lim = 12 if f in (E.F["hardship"], E.F["sale"]) else None
                texts += " ".join(str(c) for r in ws.iter_rows(max_row=lim, values_only=True) for c in r if c is not None)
    for v in ["%.1f" % x for x in g_a.values()] + ["13,178,713", "13178713", "10,583,835", "2,594,878"]:
        ok(re.search(r"(?<![\d.])%s(?![\d])" % re.escape(v), texts) is None, "no document states %s" % v)
    led = pd.read_csv(os.path.join(T, E.F["receipts"]), dtype=str)
    ok(led["receipt_ref"].is_unique and not led.duplicated().any(), "receipts: unique references, no exact duplicate rows")
    ok(int(led["loan_no"].isna().sum()) == 1, "receipts: one blank loan number")
    rec = pd.read_parquet(os.path.join(T, E.F["record"]), columns=["loan_no"])
    ok(led["loan_no"].dropna().isin(set(rec["loan_no"])).all(), "receipts: every loan number joins the record")
    for kf in ("book", "june", "defaults"):
        x = pd.read_csv(os.path.join(T, E.F[kf]), dtype=str)
        ok(x["loan_no"].is_unique and not x.duplicated().any(), "%s: one row per loan, no duplicates" % E.F[kf])
    fy, _, sale, _, _ = E.collections(T)
    ok(sale["loan_no"].is_unique, "the sale schedule lists each account once")
    ok(fy["resolution"].isin(["Cured", "Settled", "Written off"]).all() and fy["resolution_date"].notna().all(),
       "every 2023-24 case carries a resolution and a date")
    hd = pd.read_excel(os.path.join(T, E.F["hardship"]), sheet_name=0)
    ok(hd["case_ref"].is_unique and not hd.duplicated().any(), "the register: one row per case, no duplicate rows")
    sc = E.score_runs(T)
    ok(not sc.duplicated(["loan_no", "run_date"]).any(), "the scoring runs carry one row per loan and run")
    jb = pd.read_csv(os.path.join(T, E.F["june"]), dtype=str)
    ok(set(sc.loc[(sc["run_date"] > JUNE) & (sc["run_date"] < dt.date(2026, 8, 1)), "loan_no"]) <= set(jb["loan_no"]),
       "every loan of the June month end's run joins the June extract")
    pl = pd.read_csv(os.path.join(T, E.F["placements"]), dtype=str)
    ok(pl["Client Ref"].is_unique and not pl.duplicated().any(), "the placement report lists each account once, no duplicate rows")
    a2 = E.attribute(T, pd.read_csv(os.path.join(T, E.F["receipts"]), dtype={"loan_no": str, "receipt_ref": str}), "asof")
    ok(abs(led["amount_aud"].astype(float).sum() - a2["amount_aud"].astype(float).sum()) < 0.01,
       "D9: the receipts total is the same under either attribution, so no control ties only one way")
    return res


# ---------------------------------------------------------------------------------------------------------- the deliverables

QLAB = {1: ("Jul to Sep 2023", "July to September 2023"), 2: ("Oct to Dec 2023", "October to December 2023"),
        3: ("Jan to Mar 2024", "January to March 2024"), 4: ("Apr to Jun 2024", "April to June 2024")}


def deliverables(s, g_a, m_b):
    """every graded figure agrees, at its stated rounding, in the workbook, the deck and submission.md, computed here from
    target/ rather than read off golden.py; and the task tree holds one live write-up and one prompt"""
    import openpyxl
    from pptx import Presentation
    wb = openpyxl.load_workbook(os.path.join(G, "ecl_allowance_2026Q3.xlsx"))
    al = wb["Allowance"]
    ok([al["B4"].value, al["B5"].value, al["B6"].value] == [M(s["total"]), M(s["stage1"]), M(s["stage2"])],
       "workbook: the allowance and both stage amounts in AUD million to one decimal")
    ok([al["B8"].value, al["B9"].value, al["B10"].value, al["B11"].value] == [s["n1"], s["n2"], s["n3"], round(s["mult"], 2)],
       "workbook: loans by stage and the multiplier")
    lg = wb["Realised LGD"]
    got = {}
    for r in range(4, 12):
        p, ql, e = lg.cell(r, 1).value, lg.cell(r, 2).value, lg.cell(r, 5).value
        q = [k for k, v in QLAB.items() if v[0] == ql][0]
        got[(p, q)] = e
        ok(lg.cell(r, 6).value == (70 if p == "PL36" else 76), "workbook: the LGD in use beside %s %s" % (p, ql))
    ok(got == {k: round(v, 1) for k, v in g_a.items()}, "workbook: the eight realised losses at one decimal")
    mg = wb["Stage migration"]
    cells = {}
    for r in range(4, 8):
        p, s0 = mg.cell(r, 1).value, int(mg.cell(r, 2).value.split()[-1])
        for j, x in enumerate(("stage 1", "stage 2", "stage 3", "closed")):
            cells[(p, s0, x)] = mg.cell(r, 3 + j).value
    ok(cells == m_b, "workbook: the sixteen migration counts")
    prs = Presentation(os.path.join(G, "arc_ecl_q3_2026.pptx"))
    s1, s2 = prs.slides[0], prs.slides[1]
    t1 = " ".join(sh.text_frame.text for sh in s1.shapes if sh.has_text_frame)
    tb = [sh.table for sh in s1.shapes if sh.has_table][0]
    tv = [[tb.cell(i, j).text for j in range(3)] for i in range(4)]
    money = lambda x: "$%.1f million" % M(x)
    ok(money(s["total"]) in t1 and tv[1][2] == money(s["stage1"]) and tv[2][2] == money(s["stage2"])
       and tv[3][2] == money(s["total"]), "deck slide 1: the allowance and both stage amounts")
    ch = [sh.chart for sh in s2.shapes if sh.has_chart][0]
    cats = list(ch.plots[0].categories)
    order = [("PL36", 1), ("PL36", 2), ("PL60", 1), ("PL60", 2)]
    ok(cats == ["%s, stage %d at 30 June" % o for o in order], "deck slide 2: the four groups in order")
    ser = {sr.name: list(sr.values) for sr in ch.plots[0].series}
    names = {"stage 1": "Stage 1 at 30 September", "stage 2": "Stage 2", "stage 3": "Stage 3", "closed": "Closed"}
    ok(all(int(ser[names[x]][i]) == m_b[(p, s0, x)] for i, (p, s0) in enumerate(order) for x in names),
       "deck slide 2: the chart carries the sixteen counts")
    inv = {v: k for k, v in names.items()}
    ok(all(sr.points[i].data_label.has_text_frame
           and sr.points[i].data_label.text_frame.text == "{:,}".format(m_b[order[i] + (inv[sr.name],)])
           for sr in ch.plots[0].series for i in range(4)), "deck slide 2: every segment is labelled with its count")
    sub = open(os.path.join(TASK, "submission.md")).read()
    ok("AUD %.1f million" % M(s["total"]) in sub.split("## 2.")[0], "submission: the call states the allowance")
    for v in (s["total"], s["stage1"], s["stage2"]):
        ok(sub.count("AUD %.1f million" % M(v)) >= 2, "submission: AUD %.1f million in the components and the answers" % M(v))
    ok("**{:,}**".format(s["n2"]) in sub, "submission: the stage 2 count")
    for (p, q), v in g_a.items():
        ok("%s, %s: %.1f" % (p, QLAB[q][1], v) in sub, "submission: realised loss %s %s" % (p, QLAB[q][1]))
    for p, s0 in order:
        line = "%s, stage %d at 30 June: %s in stage 1, %s in stage 2, %s in stage 3, %s closed" % (
            (p, s0) + tuple("{:,}".format(m_b[(p, s0, x)]) for x in ("stage 1", "stage 2", "stage 3", "closed")))
        ok(line in sub, "submission: %s" % line)
    pr = open(os.path.join(TASK, "prompt.txt")).read().strip()
    ok(len(pr) < 1000, "the prompt is under 1,000 characters (%d)" % len(pr))
    ok(all(d in pr for d in ("ecl_allowance_2026Q3.xlsx", "arc_ecl_q3_2026.pptx")), "the prompt names both deliverables")
    for f in ("prompt.txt", "submission.md", "DESIGN_NOTE.md"):
        ok("\u2014" not in open(os.path.join(TASK, f)).read(), "no em dash in %s" % f)
    found = Counter(f for _, _, fs in os.walk(TASK) for f in fs if f in ("submission.md", "prompt.txt"))
    ok(found == Counter({"submission.md": 1, "prompt.txt": 1}), "one live write-up and one prompt in the task tree")
    ok(not any("prompt.md" in fs for _, _, fs in os.walk(TASK)), "no prompt.md beside prompt.txt")
    ok("\n" not in pr and not any(c in pr for c in "`*#|[]>"), "the prompt is one plain-text paragraph with no markdown characters")
    return len(pr)


# ---------------------------------------------------------------------------------------------------------- figures.json

def write_figures(s, out, g_a, m_b, cf, rows, jn):
    """every figure the write-up states, as it states it"""
    f = {
        "allowance_aud_million": "%.1f" % M(s["total"]), "stage1_aud_million": "%.1f" % M(s["stage1"]),
        "stage2_aud_million": "%.1f" % M(s["stage2"]), "allowance_aud": "{:,.0f}".format(s["total"]),
        "stage1_aud": "{:,.0f}".format(s["stage1"]), "stage2_aud": "{:,.0f}".format(s["stage2"]),
        "loans_stage1": "{:,}".format(s["n1"]), "loans_stage2": "{:,}".format(s["n2"]), "loans_stage3": "{:,}".format(s["n3"]),
        "loans_stages_1_2": "{:,}".format(s["n1"] + s["n2"]), "multiplier": "%.4f" % s["mult"], "multiplier_2dp": "%.2f" % s["mult"],
        "scenario_weights": ["55", "30", "15"], "scenario_multipliers": ["1.00", "1.28", "0.81"], "stage2_scalar": "2.4",
        "lgd": ["70", "76"], "past_maturity_loans": "%d" % s["n_pm"], "conversion": ["1.12", "18"], "score_fall": "60",
        "dpd_thresholds": ["30", "90", "180"], "horizon_months": "12",
        "default_path_aud_million": "%.1f" % M(out[MAIN[0][0]]["total"]),
        "default_path_aud": "{:,.0f}".format(out[MAIN[0][0]]["total"]),
        "default_path_gap_pct": "%.1f" % (100 * (1 - out[MAIN[0][0]]["total"] / s["total"])),
        "realised_loss": {"%s Q%d" % k: "%.1f" % v for k, v in sorted(g_a.items())},
        "migration": {"%s s%d %s" % k: "{:,}".format(v) for k, v in m_b.items()},
        "migration_totals": {"%s s%d" % (p, s0): "{:,}".format(sum(m_b[(p, s0, x)] for x in ("stage 1", "stage 2", "stage 3", "closed")))
                             for p in ("PL36", "PL60") for s0 in (1, 2)},
        "moved_down": "{:,}".format(sum(m_b[(p, 1, x)] for p in ("PL36", "PL60") for x in ("stage 2", "stage 3"))),
        "returned": "{:,}".format(sum(m_b[(p, 2, "stage 1")] for p in ("PL36", "PL60"))),
        "collections": {k: (("%.2f" % v) if isinstance(v, float) else "{:,}".format(v)) for k, v in cf.items()},
        "sale_price_cents": "%.2f" % W.SALE_PRICE_CENTS, "record_rows": "{:,}".format(rows["rows"]),
        "taken_over": "{:,}".format(rows["taken_over"]), "drawn_before_go_live": "{:,}".format(rows["drawn_before"]),
        "june": {"loans": "{:,}".format(jn[0]), "closed": "{:,}".format(jn[1]), "at_180": "%d" % jn[5], "at_90": "%d" % jn[6]},
        "reissue_gap_days": "90", "reporting_year": "2026",
        "policy_clauses": ["1.1", "1.2", "1.3", "1.4", "2.1", "2.2", "2.3", "3.1", "3.2", "3.3", "3.4", "3.5", "3.7", "7.1"],
        "rivals": {n: "%.1f" % M(x["total"]) for n, x in out.items()},
        "rivals_pct": {n: "%.1f" % (100 * (x["total"] / s["total"] - 1)) for n, x in out.items()},
    }
    with open(os.path.join(HERE, "figures.json"), "w") as fh:
        json.dump(f, fh, indent=1, sort_keys=True)
        fh.write("\n")


# ---------------------------------------------------------------------------------------------------------- main

def main():
    nf, nfmt, rows_ = gates()
    rows = loan_record()
    cache = os.environ.get("CHECKS_CACHE")
    if cache and os.path.exists(cache):
        import pickle
        s, out, rec = pickle.load(open(cache, "rb"))
    else:
        s, out, rec = main_call()
        if cache:
            import pickle
            pickle.dump((s, out, rec), open(cache, "wb"))
    tr = three_rules(s, out, rec)
    forks()
    realism()
    jn = june_extract()
    shares = scoring_runs()
    cf = collections_files()
    g_a, stops_a = ask_a()
    m_b, stops_b, facts, hz = ask_b()
    af = ask_forks(s, g_a, m_b)
    sep = separation(s)
    pause, stale = readings(s)
    residue_and_pins()
    au = audits(s, out, g_a, m_b)
    if T == os.path.join(TASK, "target"):
        write_figures(s, out, g_a, m_b, cf, rows, jn)
    plen = deliverables(s, g_a, m_b)
    say("CHECKS PASS", n_ok)
    say("files %d, formats %d, rows %s" % (nf, nfmt, rows_))
    say("record %s" % rows)
    say("allowance %.0f (%.1f), stage 1 %.0f (%.1f), stage 2 %.0f (%.1f); stage 2 loans %d" % (
        s["total"], M(s["total"]), s["stage1"], M(s["stage1"]), s["stage2"], M(s["stage2"]), s["n2"]))
    for name, _, kind, sf in MAIN:
        x = out[name]
        say("  %-5s %-82s %14.0f %+7.2f%%  %.1f" % (kind, name, x["total"], 100 * (x["total"] / s["total"] - 1), M(x["total"])))
    say("three rules: easy route %.0f, decisive step alone %.0f" % (tr["easy"], tr["one"]))
    say("June extract %d loans, %d closed by 30 September, %d drawn since; %d in arrears, %d on a due date counted to 30 June, "
        "%d at 180 days, %d at 90" % jn)
    say("score-fall shares by run %s" % {k: round(v, 4) for k, v in shares.items()})
    say("Collections: %s" % cf)
    keys = sorted(g_a)
    say("ask A  %-58s " % "gold" + " ".join("%6.2f" % g_a[k] for k in keys))
    for v in [n for n, *_ in ASK_A_STOPS] + [n for n, _ in ASK_A_MISREADINGS]:
        say("ask A  %-58s " % v + " ".join("%6.2f" % stops_a[v][k] for k in keys)
            + "   nearest %.2f" % min(abs(stops_a[v][k] - g_a[k]) for k in keys))
    say("ask B  gold " + " ".join("%6d" % m_b[k] for k in m_b))
    for name, _ in ASK_B_STOPS:
        x = stops_b[name]
        say("ask B  %-40s moves %2d/16" % (name, sum(x[k] != m_b[k] for k in m_b)))
    say("register at 30 June %s; September label staging moves %d counts" % (facts, hz))
    say("ask forks %s" % af)
    say("separation %s" % sep)
    say("pause dates out: %d loans, %d dates, %d main figures move, allowance %.0f; stale labels %d, bulk COVID-19 rows %d" % (pause + stale))
    say("audits %s" % au)
    say("deliverables agree with target/; prompt %d characters" % plen)


if __name__ == "__main__":
    main()
