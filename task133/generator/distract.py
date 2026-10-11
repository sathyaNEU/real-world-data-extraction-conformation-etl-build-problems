"""Files a credit risk team keeps that look relevant to the allowance and do not set it."""
import openpyxl
from openpyxl.styles import Font, PatternFill


def roll_rates_xlsx(path, rng):
    """collections MI: month-end roll rates between arrears buckets, FY26"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Roll rates FY26"
    ws.append(["Personal loans: arrears roll rates, month end, FY26 (July 2025 to June 2026)"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    hdr = ["Month", "Current to 1-29", "1-29 to 30-59", "30-59 to 60-89", "60-89 to 90+", "Cure rate 1-29"]
    ws.append(hdr)
    for j in range(1, len(hdr) + 1):
        ws.cell(row=3, column=j).font = Font(bold=True)
        ws.cell(row=3, column=j).fill = PatternFill("solid", fgColor="E5E9F0")
    for m in ["Jul-25", "Aug-25", "Sep-25", "Oct-25", "Nov-25", "Dec-25", "Jan-26", "Feb-26", "Mar-26", "Apr-26", "May-26", "Jun-26"]:
        ws.append([m, round(float(rng.uniform(0.031, 0.042)), 4), round(float(rng.uniform(0.21, 0.29)), 4),
                   round(float(rng.uniform(0.44, 0.56)), 4), round(float(rng.uniform(0.62, 0.74)), 4),
                   round(float(rng.uniform(0.52, 0.61)), 4)])
    for col, w in zip("ABCDEF", [10, 16, 15, 16, 15, 15]):
        ws.column_dimensions[col].width = w
    wb.save(path)
