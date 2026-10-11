"""The golden deliverables: the allowance workbook and the two Audit and Risk Committee slides."""
import datetime as dt
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import estimators as E
import officex as OX
import world as W
import postprocess as P

TASK = os.path.dirname(HERE)
PREPARED = dt.datetime(2026, 10, 9, 16, 20)
NAVY = RGBColor(0x1E, 0x2B, 0x3C)
TEAL = RGBColor(0x2E, 0x6F, 0x73)
SAND = RGBColor(0xC9, 0xA2, 0x5F)
GREY = RGBColor(0x5B, 0x66, 0x70)


def k(x):
    return int(round(x / 1000.0)) * 1000


def years(s, prod):
    return sorted(y for (p, y) in s["by_band"] if p == prod)


def workbook(s, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Allowance"
    ws["A1"] = "Personal loans (PL36, PL60): collective ECL allowance at 30 September 2026"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Credit Risk Reporting, 9 October 2026. CRM-ECL-04 v4.1; scenarios approved 18 September 2026."
    ws["A2"].font = Font(italic=True, color="5B6670")
    rows = [("Collective ECL allowance (AUD)", k(s["total"]), '"$"#,##0'),
            ("Stage 1 allowance (AUD)", k(s["stage1"]), '"$"#,##0'),
            ("Stage 2 allowance (AUD)", k(s["stage2"]), '"$"#,##0'),
            ("Probability-weighted PD multiplier", round(s["mult"], 4), "0.0000"),
            ("Loans in stage 3, outside the collective model", s["n3"], "#,##0")]
    for i, (lab, v, fmt) in enumerate(rows, 4):
        ws.cell(row=i, column=1, value=lab)
        c = ws.cell(row=i, column=2, value=v)
        c.number_format = fmt
    ws.cell(row=4, column=1).font = Font(bold=True)
    ws.cell(row=4, column=2).font = Font(bold=True)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 18

    ws = wb.create_sheet("Stage and product")
    hdr = ["Stage", "Product", "Loans", "Gross carrying amount (AUD)", "Allowance (AUD)", "Coverage (%)"]
    ws.append(hdr)
    for j in range(1, 7):
        c = ws.cell(row=1, column=j)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E2B3C")
        c.alignment = Alignment(wrap_text=True)
    for st in (1, 2):
        for p in ("PL36", "PL60"):
            L = s["loans"]
            g = L[(L["stage"] == st) & (L["product"] == p)]
            ws.append(["Stage %d" % st, p, int(len(g)), k(s["%s_s%d_gca" % (p, st)]), k(s["%s_s%d" % (p, st)]),
                       round(s["%s_s%d_cov" % (p, st)], 2)])
    for r in ws.iter_rows(min_row=2):
        r[2].number_format = "#,##0"
        r[3].number_format = '"$"#,##0'
        r[4].number_format = '"$"#,##0'
        r[5].number_format = "0.00"
    for col, w in zip("ABCDEF", [10, 10, 10, 26, 18, 13]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("By drawdown year")
    ws.append(["Drawdown year", "PL36 allowance (AUD)", "PL60 allowance (AUD)"])
    for j in range(1, 4):
        ws.cell(row=1, column=j).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=1, column=j).fill = PatternFill("solid", fgColor="1E2B3C")
    ys = sorted({y for (_, y) in s["by_band"]}, key=lambda y: (y != "2023 and earlier", y))
    for y in ys:
        a = s["by_band"].get(("PL36", y))
        b = s["by_band"].get(("PL60", y))
        ws.append([y, k(a) if a is not None else None, k(b) if b is not None else None])
    for r in ws.iter_rows(min_row=2):
        r[1].number_format = '"$"#,##0'
        r[2].number_format = '"$"#,##0'
    for col, w in zip("ABC", [14, 20, 20]):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet("Notes")
    notes = [
        "Scope: PL36 and PL60 loans open in the loan system at close of business 30 September 2026; loans 90 or more days past "
        "due are stage 3 and sit under CRM-ECL-06 (CRM-ECL-04 section 1).",
        "Stage 2: 30 or more days past due, a hardship arrangement in force on 30 September 2026 (start on or before, end on or "
        "after, not cancelled), or a behaviour score 60 or more points below origination on scorecard v4; v3 origination scores "
        "converted as 1.12 x v3 - 18, rounded (section 2 and the June 2024 migration note).",
        "PD(k) is (C(a+k) - C(a+k-1)) / (1 - C(a)) on the PLPD-2026.06 curve for the loan's product, where a is the number of "
        "monthly repayments fallen due by 30 September, repayment dates in a payment pause included: the curve runs from "
        "origination, and a performing loan has already come through its first a months. A loan past its contractual maturity "
        "outside stage 3 is in the final month of its term with one month remaining (section 3.2), the last month of its "
        "product's curve.",
        "EAD re-amortises the current balance over the remaining term at the contract rate; months discounted at the contract rate; "
        "LGD 70 per cent PL36, 76 per cent PL60 (sections 3.3 to 3.5).",
        "Scenario weights 55 / 30 / 15 per cent on PD multipliers 1.00 / 1.30 / 0.80 give a weighted multiplier of %.4f (section 4)." % s["mult"],
        "Amounts are rounded to the nearest thousand dollars from unrounded loan-level results, so rows may not add exactly.",
    ]
    for i, t in enumerate(notes, 1):
        ws.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 125
    # the quarter's workbook is rolled forward (save as) from the one Credit Risk Reporting set up in February 2023
    OX.xlsx_from_openpyxl(wb, path, "Joseph Mitchell", "Joseph Mitchell", dt.datetime(2023, 2, 6, 9, 31, 44), PREPARED, seed=1009)


def deck(s, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
    blank = prs.slide_layouts[6]

    def title(slide, text, sub=None):
        tb = slide.shapes.add_textbox(Cm(1.6), Cm(1.0), Cm(30.5), Cm(2.6))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = text
        r.font.size = Pt(24)
        r.font.bold = True
        r.font.color.rgb = NAVY
        if sub:
            p2 = tf.add_paragraph()
            r2 = p2.add_run()
            r2.text = sub
            r2.font.size = Pt(13)
            r2.font.color.rgb = GREY

    def footer(slide, n):
        tb = slide.shapes.add_textbox(Cm(1.6), Cm(17.7), Cm(30.5), Cm(0.8))
        r = tb.text_frame.paragraphs[0].add_run()
        r.text = "Bellbird Personal Finance  |  Audit and Risk Committee, 23 October 2026  |  Credit Risk Reporting  |  %d" % n
        r.font.size = Pt(9)
        r.font.color.rgb = GREY

    s1 = prs.slides.add_slide(blank)
    title(s1, "Book a collective ECL allowance of $%s on personal loans at 30 September 2026" % "{:,}".format(k(s["total"])),
          "Stage 1 $%s, stage 2 $%s; %s loans in stage 3 sit outside the collective model"
          % ("{:,}".format(k(s["stage1"])), "{:,}".format(k(s["stage2"])), "{:,}".format(s["n3"])))
    L = s["loans"]
    rows = [("Stage", "Product", "Loans", "Gross carrying amount", "Allowance", "Coverage")]
    for st in (1, 2):
        for p in ("PL36", "PL60"):
            g = L[(L["stage"] == st) & (L["product"] == p)]
            rows.append(("Stage %d" % st, p, "{:,}".format(len(g)), "$" + "{:,}".format(k(s["%s_s%d_gca" % (p, st)])),
                         "$" + "{:,}".format(k(s["%s_s%d" % (p, st)])), "%.2f%%" % s["%s_s%d_cov" % (p, st)]))
    tbl = s1.shapes.add_table(len(rows), 6, Cm(1.6), Cm(5.2), Cm(30.5), Cm(7.0)).table
    widths = [3.6, 3.6, 4.2, 7.4, 6.4, 5.3]
    for j, w in enumerate(widths):
        tbl.columns[j].width = Cm(w)
    from pptx.enum.text import PP_ALIGN
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = v
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (RGBColor(0xF3, 0xF5, 0xF7) if i % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF))
            para = cell.text_frame.paragraphs[0]
            para.runs[0].font.size = Pt(14)
            para.runs[0].font.bold = i == 0
            para.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else NAVY
            if j >= 2:
                para.alignment = PP_ALIGN.RIGHT
    note = s1.shapes.add_textbox(Cm(1.6), Cm(13.0), Cm(30.5), Cm(3.5)).text_frame
    note.word_wrap = True
    r = note.paragraphs[0].add_run()
    r.text = ("Probability of default is read from the PLPD-2026.06 curve from each loan's months on book at the reporting date, "
              "conditional on the loan still performing, with the 2.4 stage 2 scalar; scenarios weighted 55 / 30 / 15 per cent "
              "(multiplier %.4f). Amounts to the nearest thousand dollars." % s["mult"])
    r.font.size = Pt(11)
    r.font.color.rgb = GREY
    footer(s1, 1)

    s2 = prs.slides.add_slide(blank)
    ys = sorted({y for (_, y) in s["by_band"]}, key=lambda y: (y != "2023 and earlier", y))
    worst = max(s["by_band"].items(), key=lambda kv: kv[1])
    title(s2, "%s %s lending carries the largest share of the allowance, $%s" % (worst[0][1], worst[0][0], "{:,}".format(k(worst[1]))),
          "Collective ECL allowance by drawdown year and product, AUD, 30 September 2026")
    cd = CategoryChartData()
    cd.categories = [str(y) for y in ys]
    cd.add_series("PL36", [round(k(s["by_band"].get(("PL36", y), 0.0)) / 1e6, 3) for y in ys])
    cd.add_series("PL60", [round(k(s["by_band"].get(("PL60", y), 0.0)) / 1e6, 3) for y in ys])
    gf = s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Cm(1.6), Cm(4.6), Cm(30.5), Cm(12.5), cd)
    ch = gf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    ch.value_axis.has_major_gridlines = True
    ch.value_axis.tick_labels.number_format = '"$"0.0"m"'
    ch.value_axis.tick_labels.number_format_is_linked = False
    ch.value_axis.tick_labels.font.size = Pt(11)
    ch.category_axis.tick_labels.font.size = Pt(12)
    plot = ch.plots[0]
    for ser, col in zip(plot.series, (TEAL, SAND)):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = col
        for idx, v in enumerate(ser.values):
            if v is not None and v >= 0.1:
                dl = ser.points[idx].data_label
                dl.has_text_frame = True
                dl.text_frame.text = "$%.1fm" % v
                dl.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
                dl.text_frame.paragraphs[0].runs[0].font.color.rgb = NAVY
    footer(s2, 2)
    prs.core_properties.author = "Joseph Mitchell"
    prs.core_properties.last_modified_by = "Joseph Mitchell"
    prs.core_properties.title = "ECL allowance, Q1 FY27"
    prs.core_properties.created = PREPARED
    prs.core_properties.modified = PREPARED
    prs.core_properties.revision = 3
    prs.save(path)


def build(target=None, out=None):
    target = target or os.path.join(TASK, "target")
    out = out or os.path.join(TASK, "golden")
    os.makedirs(out, exist_ok=True)
    s = E.solve(target)
    wbp = os.path.join(out, "ecl_allowance_2026Q3.xlsx")
    workbook(s, wbp)
    dp = os.path.join(out, "arc_ecl_q3_2026.pptx")
    deck(s, dp)
    P.office_container(dp, "Joseph Mitchell", "Joseph Mitchell", dt.datetime(2026, 10, 8, 16, 41), PREPARED,
                       "Microsoft Office PowerPoint", revision=3)
    P.deterministic_embeddings(dp, dt.datetime(2026, 10, 8, 16, 41))
    for f in (wbp, dp):
        ts = W.posix(PREPARED)
        os.utime(f, (ts, ts))
    return s


if __name__ == "__main__":
    s = build()
    print("total", k(s["total"]), "stage1", k(s["stage1"]), "stage2", k(s["stage2"]), "n3", s["n3"])
