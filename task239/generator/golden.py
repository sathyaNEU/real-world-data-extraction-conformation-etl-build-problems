"""task239 golden deliverables: the structure check workbook and the one slide for the executive team."""
import os

from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import estim as E
import params as P
import postprocess
import writeup as WU

XLSX = "lfp_tci_structure_check_2027.xlsx"
PPTX = "lfp_tci_items_2028.pptx"
TEAL, ORANGE, GREY, INK = "1F6F6B", "C8553D", "8A9499", "22303A"


def _head(ws, title, sub):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13, color=INK)
    ws["A2"] = sub
    ws["A2"].font = Font(italic=True, size=9, color="555555")


def _table(ws, row, header, rows, widths, formats):
    thin = Side(style="thin", color="B7C2C6")
    for j, h in enumerate(header, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=TEAL)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(bottom=thin)
    for i, r in enumerate(rows, row + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            if formats[j - 1]:
                c.number_format = formats[j - 1]
                c.alignment = Alignment(horizontal="right", vertical="top")
            else:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(bottom=thin)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def workbook(path, A):
    wb = Workbook()
    ws = wb.active
    ws.title = "Check"
    _head(ws, "Team Climate Index structure check, 2027 survey",
          "Ten items stay for 2028; TC11 and TC12 are withdrawn (team-level loading under 0.75). %s teams, %s returns."
          % (WU.fmt(A["teams"]), WU.fmt(A["used"])))
    rows = []
    for c in E.ITEM_CODES:
        rows.append([c, WU.WORDING[c], A["loadings"][c], "Stays" if c in A["stay"] else "Withdrawn"])
    _table(ws, 4, ["Item", "Wording", "Team-level loading", "2028"], rows, [8, 70, 12, 12], [None, None, "0.00", None])
    for i, c in enumerate(E.ITEM_CODES, 5):
        if c not in A["stay"]:
            for j in range(1, 5):
                ws.cell(row=i, column=j).font = Font(bold=True, color=ORANGE)
    w2 = wb.create_sheet("Returns by channel")
    _head(w2, "2027 returns by channel", "Returns that count under WI-SV-03 section 3, down to the returns the check used.")
    rows = []
    tot = [0, 0, 0, 0]
    for ch, nm in WU.CHANNELS:
        x = A["ledger"][ch]
        v = [x["count"], x["outside"], x["incomplete"], x["used"]]
        tot = [a + b for a, b in zip(tot, v)]
        rows.append([nm[0].upper() + nm[1:]] + v)
    rows.append(["All channels"] + tot)
    _table(w2, 4, ["Channel", "Returns that count", "Outside the teams the dashboard reports",
                   "Leave an item without an answer on the scale", "Used by the check"], rows,
           [16, 14, 20, 22, 14], [None, "#,##0", "#,##0", "#,##0", "#,##0"])
    for j in range(1, 6):
        w2.cell(row=8, column=j).font = Font(bold=True)
    w3 = wb.create_sheet("Favourable answers")
    _head(w3, "Team Climate Index items: 2027 favourable answers, company-wide",
          "As the 2027 dashboard counts them: every return that counts; don't know and unanswered are not answers on "
          "the scale.")
    rows = [[c, WU.WORDING[c], A["fav"][c]["base"], A["fav"][c]["fav"], A["fav"][c]["fav"] / A["fav"][c]["base"]]
            for c in E.ITEM_CODES]
    _table(w3, 4, ["Item", "Wording", "Answers on the scale", "Favourable answers", "Percent favourable"], rows,
           [8, 64, 13, 13, 12], [None, None, "#,##0", "#,##0", "0.0%"])
    w4 = wb.create_sheet("Divisions")
    _head(w4, "Teams reported and returns used, by division",
          "Divisions as they stood on the census date, 1 September 2027.")
    rows = [[d, A["div"][d]["teams"], A["div"][d]["used"]] for d in [x[1] for x in P.DIVISIONS]]
    rows.append(["All divisions", sum(r[1] for r in rows), sum(r[2] for r in rows)])
    _table(w4, 4, ["Division", "Teams the dashboard reports", "Returns used by the check"], rows, [32, 16, 16],
           [None, "#,##0", "#,##0"])
    for j in range(1, 4):
        w4.cell(row=11, column=j).font = Font(bold=True)
    w5 = wb.create_sheet("Notes")
    notes = [
        "Standard: WI-SV-03, Staff Voice survey measurement standard, 2027 edition (sections 3 to 5).",
        "Team-level loading: one common factor fitted to the twelve items' between-team correlations; the between-team "
        "covariance is the scaled between-team matrix less the pooled within-team matrix, so the members' own covariance "
        "across the items, which every team mean carries in part, is taken out.",
        "Scores: 1 strongly disagree to 5 strongly agree; TC04 and TC07, worded negatively, scored the other way round. "
        "Paper codes 8 (more than one box marked) and 9 (no mark), like don't know, are not answers on the scale.",
        "A member who returned more than one form or session counts once, at the return received last; a tablet session "
        "the tablet sent twice counts once.",
        "Tablet returns started before the tablets' second team list load on 13 September and paper packs carry July "
        "team codes, read through HR's reorganisation map of 1 September 2027.",
        "Percent favourable: as the 2027 dashboard methodology note sets out.",
        "Data as of 19 November 2027.",
    ]
    w5["A1"] = "Notes"
    w5["A1"].font = Font(bold=True, size=12)
    for i, t in enumerate(notes, 3):
        w5.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True, vertical="top")
    w5.column_dimensions["A"].width = 110
    for sh in (ws, w2, w3, w4):
        sh.sheet_view.showGridLines = False
        sh.print_title_rows = "4:4"
    wb.save(path)
    postprocess.scrub(path, "Workforce Insight", "2027-11-19")
    postprocess.finish_office(path, "2027-11-19T09:40:00Z", "2027-11-19T15:25:00Z", "Lowfirth Parcels")


def slide(path, A):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.util import Emu, Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.6))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.text = "Ten Team Climate Index items stay for 2028; TC11 and TC12 are withdrawn"
    p.runs[0].font.size, p.runs[0].font.bold = Pt(24), True
    p.runs[0].font.color.rgb = RGBColor.from_string(INK)
    sub = s.shapes.add_textbox(Inches(0.6), Inches(1.0), Inches(12.1), Inches(0.7))
    sub.text_frame.word_wrap = True
    q = sub.text_frame.paragraphs[0]
    q.text = ("At team level TC11 and TC12 load %.2f and %.2f, nowhere near the standard's %.2f line; the other ten "
              "items, the two team-practice items included, all load well above it." % (A["loadings"]["TC11"],
                                                                                       A["loadings"]["TC12"], E.LINE))
    q.runs[0].font.size = Pt(14)
    q.runs[0].font.color.rgb = RGBColor.from_string("44525C")
    cd = CategoryChartData()
    cats = list(reversed(E.ITEM_CODES))
    cd.categories = cats
    cd.add_series("Team-level loading", [A["loadings"][c] for c in cats])
    x, y, cx, cy = Inches(0.6), Inches(1.8), Inches(8.6), Inches(5.0)
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, x, y, cx, cy, cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    ch.font.size = Pt(11)
    va = ch.value_axis
    lo, hi = -0.25, 1.0
    va.minimum_scale, va.maximum_scale, va.major_unit = lo, hi, 0.25
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor.from_string("DDE3E6")
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = "0.00", False
    plot = ch.plots[0]
    plot.gap_width = 45
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = "0.00", False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(10)
    ser = plot.series[0]
    for k, c in enumerate(cats):
        pt = ser.points[k]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = RGBColor.from_string(TEAL if c in A["stay"] else ORANGE)
    # a fixed inner plot area, so the standard's line sits at 0.75 on the value axis
    lx, ly, lw, lh = 0.09, 0.03, 0.86, 0.88
    nsC = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    pa = ch._chartSpace.find(".//{%s}plotArea" % nsC)
    lay = pa.find("{%s}layout" % nsC)
    if lay is None:
        lay = etree.SubElement(pa, "{%s}layout" % nsC)
        pa.remove(lay)
        pa.insert(0, lay)
    for child in list(lay):
        lay.remove(child)
    ml = etree.SubElement(lay, "{%s}manualLayout" % nsC)
    for tag, val in (("layoutTarget", "inner"), ("xMode", "edge"), ("yMode", "edge"), ("x", lx), ("y", ly),
                     ("w", lw), ("h", lh)):
        etree.SubElement(ml, "{%s}%s" % (nsC, tag)).set("val", str(val))
    xl = x + int(cx * (lx + lw * (E.LINE - lo) / (hi - lo)))
    line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(xl), y + int(cy * ly), Emu(xl), y + int(cy * (ly + lh)))
    line.line.color.rgb = RGBColor.from_string(INK)
    line.line.width = Pt(1.75)
    line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    lab = s.shapes.add_textbox(Emu(xl - Inches(0.9)), y - Inches(0.05), Inches(1.8), Inches(0.3))
    r = lab.text_frame.paragraphs[0]
    r.text = "Standard's line %.2f" % E.LINE
    r.runs[0].font.size, r.runs[0].font.bold = Pt(10), True
    side = s.shapes.add_textbox(Inches(9.5), Inches(2.0), Inches(3.3), Inches(4.4))
    tf = side.text_frame
    tf.word_wrap = True
    items = [("Stay for 2028", ", ".join(A["stay"]), TEAL), ("Withdraw", ", ".join(A["gone"]), ORANGE),
             ("Check", "%s teams with five or more returns, %s returns answering all twelve items"
              % (WU.fmt(A["teams"]), WU.fmt(A["used"])), INK)]
    for k, (h, t, col) in enumerate(items):
        para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        para.text = h
        para.runs[0].font.bold, para.runs[0].font.size = True, Pt(14)
        para.runs[0].font.color.rgb = RGBColor.from_string(col)
        body = tf.add_paragraph()
        body.text = t
        body.runs[0].font.size = Pt(12)
        body.space_after = Pt(10)
    foot = s.shapes.add_textbox(Inches(0.6), Inches(6.9), Inches(12.1), Inches(0.4))
    foot.text_frame.word_wrap = True
    f = foot.text_frame.paragraphs[0]
    f.text = ("Team-level loading: one factor fitted to the items' between-team correlations, the members' own "
              "covariance taken out (WI-SV-03 section 4). Source: Workforce Insight, Staff Voice 2027 structure check.")
    f.runs[0].font.size = Pt(9)
    f.runs[0].font.color.rgb = RGBColor.from_string("5B666D")
    prs.core_properties.author = "Carolyn Burrows"
    prs.core_properties.last_modified_by = "Carolyn Burrows"
    prs.core_properties.title = "Team Climate Index items for 2028"
    prs.core_properties.comments = ""
    prs.save(path)
    postprocess.scrub(path, "Lowfirth Parcels", "2027-11-19")
    postprocess.finish_office(path, "2027-11-19T11:10:00Z", "2027-11-19T16:05:00Z", "Lowfirth Parcels")


def write(folder, A):
    os.makedirs(folder, exist_ok=True)
    workbook(os.path.join(folder, XLSX), A)
    slide(os.path.join(folder, PPTX), A)
