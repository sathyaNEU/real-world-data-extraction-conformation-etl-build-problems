"""task239 documents: the company's survey measurement standard (pdf), the dashboard methodology note and the fieldwork
close-out (docx), the questionnaire codebook (xlsx), the data dictionary and the readme (md)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess

BOARD = "Lowfirth Parcels"
UNIT = "Workforce Insight"
SURVEY = "Staff Voice survey"
PEOPLE = {"head": "Carolyn Burrows", "methods": "George Haynes", "programme": "Laura Miles"}
GROUP_CODES = ["%d %s" % (k + 1, g) for k, g in enumerate(P.STAFF_GROUPS)]
YEAR_CODES = ["1 less than a year", "2 one to two years", "3 three to five years", "4 six to ten years",
              "5 more than ten years"]


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.4, leading=13, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=14, leading=18, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=14, spaceBefore=8,
                        spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10.5, textColor=colors.HexColor("#444444"))
    return dict(body=body, h1=h1, h2=h2, small=small)


def _footer(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(18 * mm, 10 * mm, text)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _table(rows, widths, size=8.2):
    t = Table(rows, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", size), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", size),
                           ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA3A8")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E3ECE8")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


def _p(text, style):
    return Paragraph(text, style)


# ------------------------------------------------------------------ the measurement standard (pdf)
def standard(path):
    S = _styles()
    b, h1, h2, sm = S["body"], S["h1"], S["h2"], S["small"]
    st = [_p("%s: measurement standard" % SURVEY, h1),
          _p("%s, %s. Standard WI-SV-03, 2027 edition, issued 3 May 2027; applies to the 2027 survey and its "
             "dashboard. Owner: %s, Head of Workforce Insight." % (BOARD, UNIT, PEOPLE["head"]), sm),
          Spacer(1, 4),
          _p("1. Purpose and scope", h2),
          _p("The %s asks every member of staff in post on the survey's census date about working for the company. "
             "This standard sets which returns count, how a return is placed in a team, how the Team Climate Index "
             "is checked each year, and which teams the dashboard reports. Other dashboard figures are made as the "
             "dashboard methodology note for the year sets out." % SURVEY, b),
          _p("2. The Team Climate Index", h2),
          _p("The index is reported for teams. It has the twelve items in Table 1, each answered on a five-point "
             "scale from strongly disagree to strongly agree, with don't know offered beside the scale. An answer on "
             "the scale is scored 1 (strongly disagree) to 5 (strongly agree), except that the two items worded "
             "negatively, TC04 and TC07, are scored the other way round, 5 for strongly disagree, so that a higher "
             "score is always the better climate.", b)]
    rows = [["Item", "Wording", "Worded"]]
    for code, wording, neg, *_ in P.ITEMS:
        rows.append([code, Paragraph(wording, ParagraphStyle("c", fontName="Helvetica", fontSize=8.2, leading=10)),
                     "negatively" if neg else "positively"])
    st += [_table(rows, [16 * mm, 128 * mm, 24 * mm]), _p("Table 1. Team Climate Index items, 2027.", sm),
           _p("3. Returns that count", h2),
           _p("3.1 An online return, whether made through a member's personal link or on a site tablet, counts when it "
              "was submitted by 23:59 on 4 October 2027. A session that was not submitted does not count.", b),
           _p("3.2 A paper form counts when %s scanned it by 8 October 2027 and it carries at least one answer." % UNIT, b),
           _p("3.3 Each return counts once, and each member's views count once: where the company received more than one "
              "return from the same member, the return received last counts and the others do not.", b),
           _p("3.4 A return's team is the team named on it: the team code loaded with the member's personal link, the "
              "team the respondent chose on a site tablet, or the team code printed on the member's paper pack. A code "
              "from a team structure that was replaced before the survey closed is read through the reorganisation map "
              "HR published with the change. A return whose code cannot be read to a single team of the structure in "
              "force on the census date, 1 September 2027, counts in the company's figures but in no team's.", b),
           _p("4. The annual structure check", h2),
           _p("4.1 Because the index reports teams, the structure check examines the twelve items at team level. A "
              "team's level on an item is a property of the team itself: the score a member of the team is expected "
              "to give, which stays the same as staff come and go.", b),
           _p("4.2 The check fits a single common factor to the twelve items' team levels across the teams in 4.4. An "
              "item's team-level loading is its loading on that factor when each item's team level is standardised "
              "across those teams.", b),
           _p("4.3 An item whose team-level loading is 0.75 or more stays in the index for the next survey. An item "
              "below 0.75 is withdrawn, and %s proposes a replacement from the item bank for the survey after." % UNIT, b),
           _p("4.4 The check uses the returns that count under section 3 from the teams reported under section 5, "
              "taking each return that gives all twelve items an answer on the scale. Don't know is not an answer on "
              "the scale.", b),
           _p("5. Reporting", h2),
           _p("5.1 A team is reported on the dashboard when it holds five or more returns that count.", b),
           _p("6. Review", h2),
           _p("%s reviews this standard after each survey. The next review is due in May 2028." % UNIT, b)]
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title="Staff Voice survey measurement standard",
                            author=PEOPLE["head"], creator=UNIT, invariant=1)
    foot = "%s | %s | WI-SV-03 (2027 edition)" % (BOARD, UNIT)
    doc.build(st, onFirstPage=_footer(foot), onLaterPages=_footer(foot))
    postprocess.scrub(path, UNIT, "2027-05-03")


# ------------------------------------------------------------------ docx helpers
def _docx(title_line, meta, paras, path, author, created, modified, title, sign=None, table=None):
    import docx
    from docx.shared import Pt
    doc = docx.Document()
    stl = doc.styles["Normal"]
    stl.font.name = "Calibri"
    stl.font.size = Pt(10.5)
    head = doc.add_paragraph()
    r = head.add_run(title_line)
    r.bold = True
    r.font.size = Pt(12)
    for k, v in meta:
        p = doc.add_paragraph()
        a = p.add_run(k + ": ")
        a.bold = True
        p.add_run(v)
        p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph()
    for t in paras:
        if isinstance(t, tuple):
            p = doc.add_paragraph()
            a = p.add_run(t[0])
            a.bold = True
            p.add_run(" " + t[1])
        else:
            p = doc.add_paragraph(t)
        p.paragraph_format.space_after = Pt(6)
    if table:
        tb = doc.add_table(rows=len(table), cols=len(table[0]))
        tb.style = "Light Grid Accent 1"
        for i, row in enumerate(table):
            for j, v in enumerate(row):
                tb.cell(i, j).text = str(v)
    if sign:
        p = doc.add_paragraph(sign)
        p.paragraph_format.space_before = Pt(8)
    doc.core_properties.author = author
    doc.core_properties.title = title
    doc.save(path)
    postprocess.scrub(path, author, created[:10])
    postprocess.finish_office(path, created, modified, BOARD)


def dashboard_note(path):
    paras = [
        ("Returns.", "The dashboard uses the returns that count under section 3 of the measurement standard "
                     "(WI-SV-03)."),
        ("Percent favourable.", "For an item, the percentage of answers on the five-point scale that are favourable: "
                                "agree or strongly agree, or, for TC04 and TC07, which are worded negatively, disagree or "
                                "strongly disagree. Don't know answers and items left unanswered are not in the base. A "
                                "figure for the company as a whole uses every return that counts, including returns whose "
                                "team is not reported and returns that leave some items unanswered."),
        ("Team pages.", "A team's page appears when the team is reported under section 5 of the standard. The page "
                        "shows each item's percent favourable for the team and the team's index score, the mean of its "
                        "twelve percentages."),
        ("Divisions.", "Each team is shown under the division it belonged to on the census date, 1 September 2027. A "
                       "move between divisions after the census date shows from the next survey."),
        ("Rounding.", "Percentages are shown to one decimal place."),
        ("Publication.", "The 2027 dashboard went live for team leaders on 1 November 2027."),
    ]
    _docx("%s | %s | 2027 dashboard" % (BOARD, UNIT),
          [("Title", "How the 2027 dashboard figures are made"), ("Author", "%s, Survey Methods Analyst" % PEOPLE["methods"]),
           ("Date", "28 October 2027")],
          paras, path, PEOPLE["methods"], "2027-10-27T14:20:00Z", "2027-10-28T09:05:00Z",
          "How the 2027 dashboard figures are made")


def close_out(path, f):
    paras = [
        "The 2027 Staff Voice survey ran from Monday 6 September to 23:59 on Monday 4 October, with paper forms taken "
        "in until Friday 8 October. Thank you to every team that made time for it.",
        ("Channels.", "Staff with a company email account got their personal link by email on the morning of 6 September. "
                      "Paper packs, each with the member's own cover letter and form, went out from 23 to 25 August to "
                      "staff without a company account and to the teams whose managers asked for paper. Site tablets were "
                      "set up at six sites in the second half of August for staff who prefer to answer at work without a "
                      "link; we ran tablet drives on the hub floors and in the depots in the first week."),
        ("Reminders.", "Email reminders went on 20 and 29 September to staff who had not yet submitted online."),
        ("Paper.", "Forms came back by internal mail and freepost and were scanned in %s on the day they arrived. A "
                   "few dozen arrived after 8 October; they were scanned all the same, but they are after "
                   "the deadline." % UNIT),
        ("Returns.", "About %s returns came in across the three channels, about %s per cent of staff in post, a little "
                     "up on 2026. The tablets were collected on 5 October." % (f["approx_returns"], f["approx_rate"])),
        ("Next.", "The dashboard is due to go live on 1 November. The structure check on the Team Climate Index follows, "
                  "in time for the 2028 questionnaire to go to the executive team in January."),
    ]
    _docx("%s | %s | Staff Voice 2027" % (BOARD, UNIT),
          [("To", "%s, Head of Workforce Insight" % PEOPLE["head"]),
           ("From", "%s, Staff Voice Programme Manager" % PEOPLE["programme"]), ("Date", "12 October 2027"),
           ("Subject", "Staff Voice 2027: fieldwork close-out")],
          paras, path, PEOPLE["programme"], "2027-10-11T15:40:00Z", "2027-10-12T08:55:00Z",
          "Staff Voice 2027 fieldwork close-out", sign=PEOPLE["programme"])


# ------------------------------------------------------------------ the codebook (xlsx)
def codebook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    ws.append(["section", "online code", "paper column", "item", "wording", "worded"])
    ws.append([1, "Q1", "Q1", "", "Which staff group do you belong to?", ""])
    ws.append([1, "Q2", "Q2", "", "How long have you worked for the company?", ""])
    for k, (code, wording, neg, *_) in enumerate(P.ITEMS):
        ws.append([2 if k < 6 else 3, "Q3%s" % "abcdefghijkl"[k], "Q3_%02d" % (k + 1), code, wording,
                   "negatively" if neg else "positively"])
    for k, (code, wording) in enumerate(P.ENGAGE):
        ws.append([4, "Q4%s" % "abcdef"[k], "Q4_%d" % (k + 1), code, wording, ""])
    w2 = wb.create_sheet("Answer codes")
    w2.append(["questions", "channel", "codes"])
    w2.append(["Q3, Q4", "online (personal link and site tablet)",
               "1 strongly disagree; 2 disagree; 3 neither agree nor disagree; 4 agree; 5 strongly agree; 6 don't know; "
               "blank not answered"])
    w2.append(["Q3, Q4", "paper form",
               "1 strongly disagree; 2 disagree; 3 neither agree nor disagree; 4 agree; 5 strongly agree; 6 don't know; "
               "8 more than one box marked; 9 no mark"])
    w2.append(["Q1", "all", "; ".join(GROUP_CODES) + "; paper also 9 no mark"])
    w2.append(["Q2", "all", "; ".join(YEAR_CODES) + "; paper also 9 no mark"])
    w3 = wb.create_sheet("Pages")
    for row in (["page", "online questions"], [1, "team; Q1, Q2"], [2, "Q3a to Q3f"], [3, "Q3g to Q3l"],
                [4, "Q4a to Q4f"], [5, "submit"]):
        w3.append(row)
    for sh, widths in ((ws, [8, 12, 12, 8, 70, 11]), (w2, [10, 34, 90]), (w3, [8, 30])):
        for c in sh[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E3ECE8")
        for i, w in enumerate(widths, 1):
            sh.column_dimensions[get_column_letter(i)].width = w
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        sh.freeze_panes = "A2"
    wb.save(path)
    postprocess.scrub(path, PEOPLE["methods"], "2027-07-14")
    postprocess.finish_office(path, "2027-07-12T10:15:00Z", "2027-07-14T13:30:00Z", BOARD)


# ------------------------------------------------------------------ the dictionary and the readme
def dictionary(F):
    return """# Data dictionary

Staff Voice 2027 files held by {unit}, {board}. Times in columns ending `_utc` are UTC; other times are UK time.

## {sessions}
One row per session the survey platform opened, for personal links and site tablets alike.
- `session_id`: the platform's number for the session, given as the session reached the platform.
- `link_type`: `personal` (the member's own emailed link) or `site` (a site tablet).
- `inv_code`: the invitation code of the personal link.
- `tablet_ref`: the tablet's own reference for the session: the tablet and the tablet's running number.
- `entered_code`: on a tablet, the invitation code the respondent typed if they chose to give it, kept as typed (the
  platform keeps the respondent's own capitals and spacing).
- `team_code`: for a personal link, the member's team code as it stood in HR's records when the links were issued on
  6 September; for a tablet, the team chosen from the tablet's list.
- `started_utc`, `submitted_utc`: when the session was started, and when the platform received the submission.
- `status`: `submitted`; `abandoned` (no activity for seven days); `open at close` (still open at 23:59 on 4 October).
- `last_page`: the last page reached; 5 is the end.

## {answers}
One row per question shown in a session, in the order shown: `session_id`, `question` (the online code in the
codebook) and `response` (the online code, blank where the question was not answered).

## {paper}
One row per form scanned. `batch` (the scanning batch), `form_serial` (the serial in the form's barcode),
`scanned_at` (when the form was scanned, on the day it arrived), then one column per question as the paper columns in
the codebook, with the paper codes.

## {packs}
One row per paper pack posted. `form_serial`; `inv_code` (the member's invitation code, printed on the pack's cover
letter); `team_code` (the team code printed on the form, from HR's team list on the day the packs were printed);
`printed`; `posted`.

## {register}
HR's team register as it stood on 15 November 2027: one row per team of the structure in force from 1 September 2027.
`team_code`, `team_name`, `division` (the team's division now), `previous_division` and `division_since` (where a team
has moved between divisions, the division it left and the date it joined its division now), `site`.

## {reorg}
HR's reorganisation map for the change of 1 September 2027: one row for each link between a team code of the July
structure and a team code of the September structure (`renamed`, `merged`, `split`), and one row for each team created
on 1 September (`new team`, no July code). Codes that did not change are not listed.

## {tablets}
The site tablets' own log: `tablet`, `site`, `event` (`team list loaded` is the tablet taking HR's team list as it
stood that day), `at_local`.

## {staff}
HR's staff in post on 1 September 2027 by team: `team_code`, `staff_in_post`, `invited` (staff in post less those on
long-term leave on 20 August, when the contact list was taken).

## {items26}
The 2026 dashboard's company-wide results for the Team Climate Index items: `item`, `wording`,
`percent_favourable_2026`, `answers_2026`.

## {codebook}
Questions, the online and paper codes, and the online pages.
""".format(unit=UNIT, board=BOARD, **F)


def readme(F):
    return """# Staff Voice 2027: survey files

{unit}'s files for the 2027 Staff Voice survey of {board}, gathered by {methods} (Survey Methods Analyst) on
Friday 19 November 2027 for the Team Climate Index structure check.

- `{standard}`: the survey's measurement standard, WI-SV-03, 2027 edition.
- `{note}`: how the 2027 dashboard figures are made.
- `{close}`: the fieldwork close-out.
- `{codebook}`: the 2027 questionnaire codebook.
- `{sessions}`, `{answers}`: the survey platform's session and response exports.
- `{paper}`: the scanned paper forms.
- `{packs}`: the paper packs posted.
- `{register}`: HR's team register, 15 November 2027.
- `{reorg}`: HR's reorganisation map, 1 September 2027.
- `{tablets}`: the site tablets' log.
- `{staff}`: HR's staff in post by team on 1 September 2027.
- `{items26}`: the 2026 dashboard's item results.
- `DATA_DICTIONARY.md`: the files and their fields.
""".format(unit=UNIT, board=BOARD, methods=PEOPLE["methods"], **F)
