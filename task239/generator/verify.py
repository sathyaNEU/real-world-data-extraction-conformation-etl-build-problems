"""Independent verifier for task239: reads only the shipped bundle, the goldens, figures.json and submission.md, and
recomputes every owed figure on its own code path (no import of the generator's estimators)."""
import json
import os
import re
import sys

import numpy as np
import pandas as pd

ITEMS = ["TC%02d" % k for k in range(1, 13)]
NEG = {"TC04", "TC07"}
ON = dict(zip(["Q3%s" % c for c in "abcdefghijkl"], ITEMS))
PA = dict(zip(["Q3_%02d" % k for k in range(1, 13)], ITEMS))


def _f(target, pat):
    return [os.path.join(target, f) for f in os.listdir(target) if re.search(pat, f)][0]


def returns(target):
    S = pd.read_csv(_f(target, r"platform_sessions"), dtype=str, keep_default_na=False)
    A = pd.read_csv(_f(target, r"platform_responses"), dtype=str, keep_default_na=False)
    Pp = pd.read_csv(_f(target, r"paper_scans"), dtype=str, keep_default_na=False)
    K = pd.read_csv(_f(target, r"paper_packs"), dtype=str, keep_default_na=False)
    S = S[S["status"] == "submitted"]
    A = A[A["question"].isin(ON)]
    wide = A.assign(item=A["question"].map(ON)).pivot(index="session_id", columns="item", values="response")
    on = S.set_index("session_id").join(wide)
    on["chan"] = on["link_type"]
    on["t"] = pd.to_datetime(on["submitted_utc"].str.rstrip("Z"))
    on["st"] = pd.to_datetime(on["started_utc"].str.rstrip("Z"))
    # a tablet session sent twice: one copy
    site = on[on["chan"] == "site"].sort_values("t")
    site = site[~site["tablet_ref"].duplicated()]
    on = pd.concat([on[on["chan"] != "site"], site])
    pk = K.set_index("form_serial")
    pa = Pp.copy()
    pa = pa[pd.to_datetime(pa["scanned_at"]).dt.date <= pd.Timestamp("2027-10-08").date()]
    qcols = [c for c in pa.columns if c.startswith("Q")]
    pa = pa[(pa[qcols] != "9").any(axis=1)]
    pa = pa.rename(columns=PA)
    pa["chan"] = "paper"
    pa["t"] = pd.to_datetime(pa["scanned_at"]) - pd.Timedelta(hours=1)
    pa["st"] = pd.NaT
    pa["team_code"] = pa["form_serial"].map(pk["team_code"])
    pa["member"] = pa["form_serial"].map(pk["inv_code"])
    known = set(S["inv_code"]) | set(K["inv_code"])
    typed = on["entered_code"].str.upper().str.replace(r"[^A-Z0-9]", "", regex=True)
    on["member"] = np.where(on["chan"] == "personal", on["inv_code"], np.where(typed.isin(known), typed, ""))
    on = on.reset_index().rename(columns={"session_id": "rid"})
    pa["rid"] = "F" + pa["form_serial"]
    cols = ["rid", "chan", "t", "st", "team_code", "member"] + ITEMS
    R = pd.concat([on[cols], pa[cols]], ignore_index=True)
    for it in ITEMS:
        v = pd.to_numeric(R[it], errors="coerce")
        v = v.where((v >= 1) & (v <= 5))
        R[it] = 6 - v if it in NEG else v
    R = R.sort_values(["t", "rid"])
    keyed = R[R["member"] != ""]
    keyed = keyed[~keyed["member"].duplicated(keep="last")]
    return pd.concat([R[R["member"] == ""], keyed]).sort_values(["t", "rid"]).reset_index(drop=True)


def teams(target, R):
    reg = pd.read_csv(_f(target, r"team_register"), dtype=str, keep_default_na=False)
    mp = pd.read_csv(_f(target, r"reorganisation_map"), dtype=str, keep_default_na=False)
    tl = pd.read_csv(_f(target, r"site_tablet_log"), dtype=str, keep_default_na=False)
    codes = set(reg["team_code"])
    succ = mp[mp["old_code"] != ""].groupby("old_code")["new_code"].apply(lambda s: sorted(set(s)))
    switch = pd.to_datetime(tl.loc[tl["event"] == "team list loaded", "at_local"]).max() - pd.Timedelta(hours=1)

    def july(c):
        if c in succ.index:
            return succ[c][0] if len(succ[c]) == 1 else None
        return c if c in codes else None

    out = []
    for ch, c, st in zip(R["chan"], R["team_code"], R["st"]):
        if ch == "personal":
            out.append(c if c in codes else None)
        elif ch == "site" and st < switch:
            out.append(july(c))
        elif ch == "site":
            out.append(c if c in codes else None)
        else:
            out.append(july(c))
    R = R.copy()
    R["team"] = out
    n = R.dropna(subset=["team"]).groupby("team").size()
    rep = set(n[n >= 5].index)
    R["reported"] = R["team"].isin(rep)
    R["complete"] = R[ITEMS].notna().all(axis=1)
    R["used"] = R["reported"] & R["complete"]
    return R, rep, reg


def between_loadings(U):
    """the scaled between-team matrix less the pooled within-team matrix, then principal axis factoring"""
    X = U[ITEMS].astype(float)
    grp = X.groupby(U["team"].values)
    n = grp.size().values.astype(float)
    M = grp.mean().values
    N, G = len(X), len(n)
    Dw = X.values - grp.transform("mean").values
    Spw = Dw.T @ Dw / (N - G)
    Dm = M - X.values.mean(0)
    Sb = (Dm * n[:, None]).T @ Dm / (G - 1)
    c = (N * N - (n * n).sum()) / (N * (G - 1))
    Sig = (Sb - Spw) / c
    d = np.sqrt(np.diag(Sig))
    Rm = Sig / np.outer(d, d)
    h = 1 - 1 / np.diag(np.linalg.inv(Rm))
    for _ in range(20000):
        Rr = Rm.copy()
        np.fill_diagonal(Rr, h)
        w, v = np.linalg.eigh(Rr)
        lam = v[:, -1] * np.sqrt(w[-1])
        if np.max(np.abs(lam ** 2 - h)) < 1e-13:
            break
        h = lam ** 2
    return lam if lam.sum() > 0 else -lam


def main(target, golden, figures, submission):
    F = json.load(open(figures))
    sub = open(submission).read()
    R = returns(target)
    R, rep, reg = teams(target, R)
    U = R[R["used"]]
    lam = between_loadings(U)
    stay = [it for it, l in zip(ITEMS, lam) if l >= 0.75]
    assert stay == F["call"]["stay"], (stay, F["call"]["stay"])
    assert all(lam[i] >= F["components"]["kept_band"] for i, it in enumerate(ITEMS) if it in stay)
    assert all(abs(lam[i]) < F["components"]["withdrawn_band"] for i, it in enumerate(ITEMS) if it not in stay)
    for i, it in enumerate(ITEMS):
        assert abs(round(float(lam[i]), 2) - F["deliverables"]["loadings"][it]) < 0.006, it
    assert len(R) == F["steps"]["returns_that_count"] and len(rep) == F["steps"]["teams_reported"]
    assert len(U) == F["steps"]["returns_used"]
    for ch in ("personal", "site", "paper"):
        x = R[R["chan"] == ch]
        mine = dict(count=len(x), outside=int((~x["reported"]).sum()),
                    incomplete=int((x["reported"] & ~x["complete"]).sum()), used=int(x["used"].sum()))
        assert mine == F["deliverables"]["ledger"][ch], (ch, mine)
    for it in ITEMS:
        v = R[it].dropna()
        assert dict(base=len(v), fav=int((v >= 4).sum())) == F["deliverables"]["favourable"][it], it
    names = {"Hub Operations", "Linehaul and Fleet", "Customer Service", "Engineering", "Delivery Operations",
             "Corporate Services"}
    regi = reg.set_index("team_code")
    div = {}
    for t in rep:
        r = regi.loc[t]
        dname = r["previous_division"] if (r["division_since"] and r["division_since"] > "2027-09-01") else r["division"]
        assert dname in names
        div.setdefault(dname, [0, 0])[0] += 1
    for t, k in U.groupby("team").size().items():
        r = regi.loc[t]
        dname = r["previous_division"] if (r["division_since"] and r["division_since"] > "2027-09-01") else r["division"]
        div[dname][1] += int(k)
    for dname, (tn, un) in div.items():
        assert F["deliverables"]["divisions"][dname] == {"teams": tn, "used": un}, dname
    # the goldens carry the same figures
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(golden, "lfp_tci_structure_check_2027.xlsx"), data_only=True)
    ws = wb["Check"]
    rows = {ws.cell(row=r, column=1).value: (ws.cell(row=r, column=3).value, ws.cell(row=r, column=4).value)
            for r in range(5, 17)}
    for i, it in enumerate(ITEMS):
        assert abs(rows[it][0] - round(float(lam[i]), 2)) < 1e-9 and rows[it][1] == ("Stays" if it in stay else "Withdrawn")
    ws = wb["Returns by channel"]
    for r, ch in zip(range(5, 8), ("personal", "site", "paper")):
        got = [ws.cell(row=r, column=c).value for c in range(2, 6)]
        x = F["deliverables"]["ledger"][ch]
        assert got == [x["count"], x["outside"], x["incomplete"], x["used"]], ch
    ws = wb["Favourable answers"]
    for r in range(5, 17):
        it = ws.cell(row=r, column=1).value
        assert [ws.cell(row=r, column=3).value, ws.cell(row=r, column=4).value] == [
            F["deliverables"]["favourable"][it]["base"], F["deliverables"]["favourable"][it]["fav"]], it
    ws = wb["Divisions"]
    for r in range(5, 11):
        dname = ws.cell(row=r, column=1).value
        assert [ws.cell(row=r, column=2).value, ws.cell(row=r, column=3).value] == [
            F["deliverables"]["divisions"][dname]["teams"], F["deliverables"]["divisions"][dname]["used"]], dname
    from pptx import Presentation
    prs = Presentation(os.path.join(golden, "lfp_tci_items_2028.pptx"))
    ch = [s for s in prs.slides[0].shapes if s.has_chart][0].chart
    vals = dict(zip(ch.plots[0].categories, ch.plots[0].series[0].values))
    for i, it in enumerate(ITEMS):
        assert abs(vals[it] - round(float(lam[i]), 2)) < 1e-9, it
    # every number in the submission's answer blocks recomputes
    for x in F["deliverables"]["ledger"].values():
        for v in x.values():
            assert "{:,}".format(v) in sub
    for x in F["deliverables"]["favourable"].values():
        assert "{:,}".format(x["base"]) in sub and "{:,}".format(x["fav"]) in sub
    print("verify: every figure recomputes from the bundle (%d returns, %d teams, %d used; %s stay)"
          % (len(R), len(rep), len(U), len(stay)))


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    task = os.path.dirname(here)
    main(os.path.join(task, "target"), os.path.join(task, "golden"), os.path.join(here, "figures.json"),
         os.path.join(task, "submission.md"))
