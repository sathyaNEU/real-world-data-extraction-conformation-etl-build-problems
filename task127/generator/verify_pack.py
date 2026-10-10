#!/usr/bin/env python3
"""task127 independent verifier: reads only the shipped folder (target/) and recomputes every rung, the
rival-killers, the calibration back-test and every graded figure, on its own code path (csv, json and
openpyxl; no generator module, no pandas).

    python3 task127/generator/verify_pack.py <task root or target dir> [--json out.json]
"""
import csv
import json
import math
import os
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl

EXPECT = {
    "slots": {"North Shore": 160, "Valley": 180, "Lakes": 240, "Uplands": 170, "Riverbend": 340, "Pinewood": 80},
    "placed": 1170, "unallocated": 630, "leader": "Riverbend", "lead_by": 100,
    "rung_leaders": ["Valley", "Uplands", "North Shore", "Lakes", "Riverbend"],
    "rung3_split": {"North Shore": 160, "Valley": 180, "Lakes": 420, "Uplands": 300, "Riverbend": 340, "Pinewood": 80},
}
COOPS = ["North Shore", "Valley", "Lakes", "Uplands", "Riverbend", "Pinewood"]
QUAL = {"PF": "furnace", "ER": "resistance", "PB": "boiler"}
PILOT_SYS = {"propane furnace (ducted)": "furnace", "electric resistance": "resistance",
             "propane boiler (hydronic)": "boiler"}
IN_BAND = ("35,000 to 74,999", "75,000 to 149,999")
CT = ZoneInfo("America/Chicago")
FAILS = []
COUNT = [0]


def check(cond, msg):
    COUNT[0] += 1
    if not cond:
        FAILS.append(msg)
        print("FAIL", msg)


def rows_csv(p):
    with open(p, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def rows_xlsx(p, sheet, header_row=1):
    ws = openpyxl.load_workbook(p, read_only=True, data_only=True)[sheet]
    it = ws.iter_rows(values_only=True)
    for _ in range(header_row - 1):
        next(it)
    hdr = [str(h) if h is not None else "" for h in next(it)]
    out = []
    for r in it:
        if r is None or all(v is None for v in r):
            continue
        out.append(dict(zip(hdr, r)))
    return out


def d10(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return date.fromisoformat(str(v)[:10])


# ------------------------------------------------------------------ households, rates, rungs

def survey(T):
    sv = rows_csv(T / "heat_survey_2025_households.csv")
    for r in sv:
        r["weight"] = int(r["weight"])
    return sv


def is_qual(r):
    return (r["tenure"] == "owner" and r["structure"] == "single-family detached" and r["income_band"] in IN_BAND
            and r["heating_system"] in QUAL)


def rungs(T, sv, pilot):
    pnb = {r["neighbourhood"] for r in pilot}
    den, num = defaultdict(int), defaultdict(int)
    for r in sv:
        if is_qual(r) and r["neighbourhood"] in pnb:
            den[QUAL[r["heating_system"]]] += r["weight"]
    for r in pilot:
        num[PILOT_SYS[r["heating_system_replaced"]]] += 1
    rate = {s: num[s] / den[s] for s in den}
    pooled = len(pilot) / sum(den.values())
    # back-test per pilot neighbourhood
    nb_q = defaultdict(lambda: defaultdict(int))
    for r in sv:
        if is_qual(r) and r["neighbourhood"] in pnb:
            nb_q[r["neighbourhood"]][QUAL[r["heating_system"]]] += r["weight"]
    act = defaultdict(int)
    for r in pilot:
        act[r["neighbourhood"]] += 1
    bt_sys = {nb: sum(q[s] * rate[s] for s in q) / act[nb] - 1 for nb, q in nb_q.items()}
    bt_pool = {nb: sum(q.values()) * pooled / act[nb] - 1 for nb, q in nb_q.items()}
    check(len(bt_sys) == 6 and max(abs(v) for v in bt_sys.values()) <= 0.04, "system rates reproduce 6 of 6")
    check(min(abs(v) for v in bt_pool.values()) >= 0.15, "pooled rate misses every pilot neighbourhood by 15%+")
    # rung 0 from the published tables
    wb = openpyxl.load_workbook(T / "heat_survey_2025_tables.xlsx", read_only=True, data_only=True)

    def table(name):
        it = wb[name].iter_rows(values_only=True)
        for _ in range(3):
            next(it)
        hdr = list(next(it))
        return [dict(zip(hdr, r)) for r in it if r and r[0]]
    h1, h2, h3, h4 = table("H1"), table("H2"), table("H3"), table("H4")
    r0 = {}
    for c in COOPS:
        a = next(r for r in h1 if r["Co-operative"] == c)
        hh = sum(v for k, v in a.items() if k != "Co-operative")
        pe = a["bottled, tank or LP gas"] + a["electricity"]
        own = sum(v for r in h2 if r["Co-operative"] == c and r["Tenure"] == "owner"
                  for k, v in r.items() if k not in ("Co-operative", "Tenure"))
        det = sum(r["single-family detached"] for r in h3 if r["Co-operative"] == c)
        b = next(r for r in h4 if r["Co-operative"] == c)
        inc = b[IN_BAND[0]] + b[IN_BAND[1]]
        r0[c] = hh * (pe / hh) * (own / hh) * (det / hh) * (inc / hh) * pooled
        tot = sum(r["weight"] for r in sv if r["coop"] == c and r["heating_system"] in ("PF", "PB"))
        check(tot == a["bottled, tank or LP gas"], f"H1 propane count reproduces for {c}")
    # rungs 1 and 2 from the joint counts, outside the pilot neighbourhoods
    r1, buyers = defaultdict(float), defaultdict(float)
    for r in sv:
        if is_qual(r) and r["neighbourhood"] not in pnb:
            r1[r["coop"]] += r["weight"] * pooled
            buyers[(r["coop"], r["neighbourhood"])] += r["weight"] * rate[QUAL[r["heating_system"]]]
    r2 = defaultdict(float)
    for (c, nb), v in buyers.items():
        r2[c] += v
    # rung 3, the feeder cap
    fmap = {(r["coop"], r["neighbourhood"]): r["feeder"] for r in rows_csv(T / "area_feeder_map.csv")}
    host = {r["feeder"]: int(r["hosting_capacity_remaining_kw"]) // 5
            for r in rows_xlsx(T / "hosting_capacity_nov2026.xlsx", "feeders", header_row=5)}
    per = defaultdict(float)
    for k, v in buyers.items():
        per[(k[0], fmap[k])] += v
    r3 = defaultdict(float)
    for (c, f), v in per.items():
        r3[c] += min(v, host[f])
    return dict(rate=rate, pooled=pooled, r0=r0, r1=dict(r1), r2=dict(r2), r3=dict(r3), buyers=buyers,
                fmap=fmap, host=host, bt_pool=bt_pool)


def share_tens(vals, total=1800):
    s = sum(vals.values())
    q = {k: v / s * total / 10 for k, v in vals.items()}
    base = {k: int(math.floor(v)) for k, v in q.items()}
    for k in sorted(q, key=lambda k: q[k] - base[k], reverse=True)[:total // 10 - sum(base.values())]:
        base[k] += 1
    return {k: 10 * v for k, v in base.items()}


def tens(v):
    return int(math.floor(v / 10 + 0.5)) * 10


# ------------------------------------------------------------------ crews and the 2027 meter sets

def us_holidays(y):
    def nth(m, wd, n):
        d = date(y, m, 1)
        d += timedelta(days=(wd - d.weekday()) % 7)
        return d + timedelta(weeks=n - 1)

    def obs(d):
        return d + timedelta(days={5: -1, 6: 1}.get(d.weekday(), 0))
    last_mon_may = date(y, 5, 31) - timedelta(days=date(y, 5, 31).weekday())
    tg = nth(11, 3, 4)
    return {obs(date(y, 1, 1)), last_mon_may, obs(date(y, 7, 4)), nth(9, 0, 1), tg, tg + timedelta(days=1),
            obs(date(y, 12, 25))}


def crew(fo, coop):
    tot, std = defaultdict(int), defaultdict(int)
    batch = []
    for r in fo:
        if r["coop"] != coop or not r["completed_date"]:
            continue
        d = r["completed_date"]
        tot[d] += 1
        if r["order_type"] == "MXCH" and r["requested_date"] == "2025-07-07":
            batch.append(d)
        elif r["order_type"] != "HPRM":
            std[d] += 1
    days = sorted(d for d in tot if "2025-01-01" <= d <= "2026-12-10")
    plateau = [d for d in days if min(batch) <= d < max(batch)]
    ceiling = sum(tot[d] for d in plateau) / len(plateau)
    win = [d for d in days if d >= "2025-12-12"]
    standing = sum(std[d] for d in win) / len(win)
    all_std = (sum(std[d] for d in days) + len(batch)) / len(days)
    return dict(ceiling=ceiling, standing=standing, spare=ceiling - standing, plateau=len(plateau), days=days,
                tot=tot, std=std, batch=len(batch), standing_with_batch=all_std)


def meter_sets(placeable, spare, arrivals, workdays):
    """Day-by-day: orders raised on a day are worked from the crew's next working day."""
    waiting, done = 0.0, 0.0
    d = date(2027, 1, 1)
    while d.year == 2027:
        if d in workdays:
            s = min(waiting, spare)
            waiting -= s
            done += s
        waiting += placeable * arrivals.get(d, 0.0)
        d += timedelta(days=1)
    return done


def queue(T, R, pilot):
    fo = rows_csv(T / "field_orders_2025_2026.csv")
    four = {}
    for c in COOPS:
        wd = {datetime.strptime(r["completed_date"], "%Y-%m-%d").weekday() for r in fo
              if r["coop"] == c and r["completed_date"]}
        if max(wd) == 3:
            four[c] = crew(fo, c)
    check(sorted(four) == ["Lakes", "Uplands"], f"four-day crews {sorted(four)}")
    for c, k in four.items():
        for h in ("2025-09-01", "2025-11-27", "2025-12-25", "2026-09-07"):
            check(k["tot"].get(h, 0) == 0, f"{c} crew off on {h}")
        check(k["tot"].get("2026-07-02", 0) > 0, f"{c} crew worked Thursday 2 July 2026")
        check(k["plateau"] >= 40, f"{c} ceiling held {k['plateau']} working days")
    arr = defaultdict(float)
    for r in pilot:
        d = d10(r["install_date"])
        arr[date(2027, d.month, d.day)] += 1 / len(pilot)
    hol = us_holidays(2027)
    wd27 = {date(2027, 1, 1) + timedelta(days=i) for i in range(365)}
    wd27 = {d for d in wd27 if d.weekday() < 4 and d not in hol}
    sets = {}
    for c in COOPS:
        sets[c] = meter_sets(R["r3"][c], four[c]["spare"], arr, wd27) if c in four else R["r3"][c]
    # closed form: a quarter in spring, then the crew's daily figure from Monday 27 September
    for c in four:
        n = len([d for d in wd27 if d >= date(2027, 9, 27)])
        check(abs(R["r3"][c] / 4 + n * four[c]["spare"] - sets[c]) < 1e-6, f"{c} closed form agrees")
    return sets, four, arr, wd27


# ------------------------------------------------------------------ the asks

def asks(T, pilot):
    from pypdf import PdfReader
    import re
    terms = " ".join(p.extract_text() for p in PdfReader(str(T / "cooperative_participation_terms.pdf")).pages)
    sched = {k: int(re.search(re.escape(lbl) + r"\s*\$([\d,]+)", terms, re.I).group(1).replace(",", ""))
             for k, lbl in (("furnace", "Propane furnace (ducted)"), ("resistance", "Electric resistance"),
                            ("boiler", "Propane boiler (hydronic)"))}
    top = int(re.search(r"plus \$([\d,]+)", terms).group(1).replace(",", ""))
    by_prem = {str(r["premises_id"]): r for r in pilot}
    by_reb = {str(r["rebate_id"]): r for r in pilot}
    inv = rows_csv(T / "installer_invoices_2026.csv")
    jobs = defaultdict(list)
    for r in inv:
        jobs[r["invoice_no"]].append(r)
    cost, n, reb = defaultdict(float), defaultdict(int), defaultdict(float)
    for no, lines in jobs.items():
        v = max(int(r["version"]) for r in lines if r["accepted_at"])
        ls = [r for r in lines if int(r["version"]) == v]
        heads = [r for r in ls if r["description"].startswith("multi-zone heat pump system, indoor head")]
        t = sum(float(r["amount_usd"]) for r in ls if r not in heads) + (float(heads[0]["amount_usd"]) if heads else 0)
        p = by_prem[ls[0]["premises_id"]]
        cost[p["coop"]] += t
        n[p["coop"]] += 1
        reb[p["coop"]] += sched[PILOT_SYS[p["heating_system_replaced"]]] + (top if p["income_band"] == IN_BAND[0] else 0)
    A = {c: (cost[c] / n[c], reb[c] / cost[c] * 100) for c in COOPS}
    check(sum(n.values()) == 412, "one invoice of record per pilot install")
    led = rows_csv(T / "rebate_payments_2026.csv")
    ret = {x["original_payment_reference"] for x in json.loads((T / "bank_returns_2026.json").read_text())["notices"]}
    paid, cov = defaultdict(float), defaultdict(set)
    lazy = defaultdict(float)
    for r in led:
        if not r["cleared_at"]:
            continue
        t = datetime.strptime(r["cleared_at"], "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=timezone.utc)
        c = by_reb[r["rebate_id"]]["coop"]
        if t.date() <= date(2026, 11, 30):
            lazy[c] += float(r["amount_usd"])
        if r["payment_id"] in ret or t.astimezone(CT).date() > date(2026, 11, 30):
            continue
        paid[c] += float(r["amount_usd"])
        cov[c].add(r["rebate_id"])
    B = {c: (paid[c], len(cov[c])) for c in COOPS}
    for c in COOPS:
        check(abs(lazy[c] - paid[c]) >= 100, f"{c} lazy ledger read misses the paid dollars")
    return A, B


EXPECT_ASKS = {   # the answer key, as recorded in the design note's build record
    "ask_a": {"North Shore": [16962, 23.6], "Valley": [16952, 23.8], "Lakes": [14786, 29.7], "Uplands": [15062, 29.2],
              "Riverbend": [16794, 24.0], "Pinewood": [14263, 31.7]},
    "ask_b": {"North Shore": [326400, 82], "Valley": [295200, 74], "Lakes": [143800, 32], "Uplands": [121000, 27],
              "Riverbend": [290800, 72], "Pinewood": [219000, 49]},
}


def main():
    args = sys.argv[1:]
    root = Path(args[0]).resolve()
    T = root / "target" if (root / "target").is_dir() else root
    out_json = args[args.index("--json") + 1] if "--json" in args else None
    files = sorted(os.listdir(T))
    check(len(files) >= 10, "ten or more files")
    check(len({Path(f).suffix for f in files}) >= 3, "three or more formats")
    nrows = sum(1 for _ in open(T / "field_orders_2025_2026.csv", encoding="utf-8")) - 1
    check(nrows >= 25000, f"field orders {nrows} rows")
    meta = T.parent / "metadata.json"
    if meta.exists():
        m = json.loads(meta.read_text())
        check(len(m["distractor_files"]) >= 2 and all(f in files for f in m["distractor_files"]),
              "two distractors named in metadata.json and shipped")
    pilot = rows_xlsx(T / "pilot_rebates_2026.xlsx", "rebates")
    for r in pilot:
        r["premises_id"], r["rebate_id"] = str(r["premises_id"]), str(r["rebate_id"])
    check(len(pilot) == 412, "412 pilot installs")
    sv = survey(T)
    R = rungs(T, sv, pilot)
    sets, four, arr, wd27 = queue(T, R, pilot)
    lead = []
    for k, v in enumerate([R["r0"], R["r1"], R["r2"], R["r3"], sets]):
        o = sorted(v, key=v.get, reverse=True)
        lead.append(o[0])
        check(v[o[0]] / v[o[1]] >= 1.15, f"rung {k} margin {v[o[0]] / v[o[1]]:.3f}")
    check(lead == EXPECT["rung_leaders"], f"rung leaders {lead}")
    check(sorted(R["r0"], key=R["r0"].get, reverse=True).index("Riverbend") + 1 in (4, 5), "answer 4th or 5th at rung 0")
    r3s = {c: tens(v) for c, v in R["r3"].items()}
    check(r3s == EXPECT["rung3_split"] and sum(r3s.values()) == 1480, f"rung 3 split {r3s}")
    sl = {c: tens(v) for c, v in sets.items()}
    placed = sum(sl.values())
    o = sorted(sl, key=sl.get, reverse=True)
    check(sl == EXPECT["slots"], f"golden split {sl}")
    check(placed == EXPECT["placed"] and 1800 - placed == EXPECT["unallocated"], f"placed {placed}")
    check(o[0] == EXPECT["leader"] and sl[o[0]] - sl[o[1]] == EXPECT["lead_by"], f"leader {o[0]}")
    for c, v in sets.items():
        r = (v - 5) % 10
        check(min(r, 10 - r) >= 1.5, f"{c} {v:.2f} clear of a bin edge")
    for k in (0, 1, 2):
        v = [R["r0"], R["r1"], R["r2"]][k]
        check(sum(v.values()) > 1800, f"rung {k} shares the 1,800")
    # rival readings: no lapse, standing ignored, batch carried as standing work
    check({c: tens(v) for c, v in R["r3"].items()} != sl, "the unqueued split is not the answer")
    full = {c: meter_sets(R["r3"][c], four[c]["ceiling"], arr, wd27) for c in four}
    check(all(abs(full[c] - R["r3"][c]) < 1e-6 for c in four), "against the full ceiling every meter is set")
    gsets = {c: meter_sets(R["r3"][c], four[c]["ceiling"] - four[c]["standing_with_batch"], arr, wd27) for c in four}
    check(all(tens(gsets[c]) < sl[c] for c in four), "batch carried as run-rate lands below the answer")
    yr = {c: len([d for d in wd27]) * four[c]["spare"] for c in four}
    check(all(yr[c] > R["r3"][c] for c in four), "each binding crew's year has room for its meters")
    A, B = asks(T, pilot)
    for c in COOPS:
        ra = (A[c][0] - 0.5) % 1
        check(min(ra, 1 - ra) >= 0.2, f"{c} cost clear of the edge")
        rs = (A[c][1] - 0.05) % 0.1
        check(min(rs, 0.1 - rs) >= 0.02, f"{c} share clear of the edge")
    res = {"slots": sl, "placed": placed, "unallocated": 1800 - placed, "leader": o[0],
           "sets": {c: round(v, 3) for c, v in sets.items()}, "rung3": {c: round(v, 3) for c, v in R["r3"].items()},
           "rates": {k: round(v * 100, 4) for k, v in R["rate"].items()}, "pooled": round(R["pooled"] * 100, 4),
           "spare": {c: round(k["spare"], 4) for c, k in four.items()},
           "ceiling": {c: round(k["ceiling"], 4) for c, k in four.items()},
           "standing": {c: round(k["standing"], 4) for c, k in four.items()},
           "ask_a": {c: [round(A[c][0]), round(A[c][1], 1)] for c in COOPS},
           "ask_b": {c: [round(B[c][0]), B[c][1]] for c in COOPS}}
    if EXPECT_ASKS:
        check(res["ask_a"] == EXPECT_ASKS["ask_a"] and res["ask_b"] == EXPECT_ASKS["ask_b"], "ask answers match the key")
    res["checks"] = COUNT[0]
    res["failures"] = FAILS
    if out_json:
        Path(out_json).write_text(json.dumps(res, indent=1))
    print(json.dumps({k: res[k] for k in ("slots", "placed", "leader", "ask_a", "ask_b")}))
    print(f"VERIFY {'OK' if not FAILS else 'FAILED'}: {COUNT[0]} checks, {len(FAILS)} failures")
    sys.exit(1 if FAILS else 0)


if __name__ == "__main__":
    main()
