"""The pack's documents: the reserving basis, the two systems' notes, the legacy code list, the ledger notes, the
close calendar and finance's preliminary reserve register."""
from __future__ import annotations

import datetime as dt

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                   # noqa: E402
from reportlab.lib.pagesizes import LETTER                                         # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                    # noqa: E402
from reportlab.lib.units import inch                                               # noqa: E402
from reportlab.pdfbase import pdfmetrics                                           # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                       # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer                # noqa: E402

import common as K   # noqa: E402
import world as W    # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("Serif", "LiberationSerif-Regular"), ("Serif-B", "LiberationSerif-Bold"), ("Sans", "LiberationSans-Regular")):
    pdfmetrics.registerFont(TTFont(nm, FONT_DIR + f + ".ttf"))
INK, GREY, GREEN = colors.HexColor("#1b1f1d"), colors.HexColor("#5d6661"), colors.HexColor("#2f5d3a")
BODY = ParagraphStyle("b", fontName="Serif", fontSize=10.5, leading=14.2, textColor=INK, spaceAfter=6)
CL = ParagraphStyle("c", parent=BODY, leftIndent=26, firstLineIndent=-26)
H1 = ParagraphStyle("h1", fontName="Serif-B", fontSize=15, leading=19, textColor=GREEN, spaceAfter=3)
H2 = ParagraphStyle("h2", fontName="Serif-B", fontSize=11, leading=14.5, textColor=INK, spaceBefore=10, spaceAfter=4)
META = ParagraphStyle("m", fontName="Sans", fontSize=8.6, leading=11.5, textColor=GREY, spaceAfter=12)


def _frame(left, right):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(GREY)
        canvas.setFont("Sans", 7.8)
        canvas.drawString(1.0 * inch, LETTER[1] - 0.6 * inch, left)
        canvas.drawRightString(LETTER[0] - 1.0 * inch, LETTER[1] - 0.6 * inch, right)
        canvas.drawRightString(LETTER[0] - 1.0 * inch, 0.55 * inch, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def write_basis(path):
    act = K.PEOPLE["actuary"]
    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=1.0 * inch, rightMargin=1.0 * inch,
                            topMargin=0.95 * inch, bottomMargin=0.9 * inch,
                            title="Reserving basis: private passenger auto liability, bodily injury", author=act,
                            subject="", creator="", producer="")
    story = [Paragraph("Reserving basis: private passenger auto liability, bodily injury", H1),
             Paragraph("%s, Actuarial Department. Version 2026.1, issued 10 November 2026 by %s, Chief Actuary; "
                       "replaces version 2023.2." % (K.ORG, act), META)]
    sections = [
        ("1. Scope", [
            "1.1 This basis governs the actuarial central estimate of unpaid loss for private passenger auto "
            "liability, bodily injury, on the company's business in Iowa, Nebraska and South Dakota. Other lines and "
            "coverages have their own bases.",
            "1.2 The basis is reviewed each year before the year-end valuation."]),
        ("2. The estimate", [
            "2.1 The actuarial central estimate at a valuation date is the amount of loss payments still to be made "
            "after that date on bodily injury claimant features from accidents on or before it, including claims not "
            "yet reported to the company.",
            "2.2 It excludes adjusting expenses (defence and cost containment, and adjusting and other), which are "
            "estimated separately. It is gross of reinsurance and undiscounted. Recoveries do not arise on this "
            "coverage.",
            "2.3 The selection of methods and assumptions is the actuary's at each review and is documented in it."]),
        ("3. Data conventions", [
            "3.1 The count unit is the claimant feature: one injured person's bodily injury claim on one accident.",
            "3.2 Accident quarters are the calendar quarters of the accident date. Development is measured at quarter "
            "ends: age 1 is the end of the accident quarter itself, age 2 the end of the following quarter, and so on.",
            "3.3 Paid loss is net of voided checks, each void on its own date.",
            "3.4 The company's books, and every valuation date in this basis, run on Central time.",
            "3.5 Bodily injury features settle within forty quarters of the end of their accident quarter; the "
            "estimate carries no development beyond age 40.",
            "3.6 Case estimates are not used for bodily injury: the estimate rests on paid loss and claimant feature "
            "counts."]),
        ("4. Data sources", [
            "4.1 Claims data come from the legacy claims system (its feature archive and payment ledger, for work "
            "done before the conversion weekend of 29 and 30 June 2024) and from the claims platform (its feature "
            "register, feature events and payments, for work done since 1 July 2024).",
            "4.2 Paid loss is tied each year to general ledger account 5311, losses paid, auto bodily injury."]),
        ("5. Booking", [
            "5.1 Finance books a preliminary reserve at each quarter end from its roll-forward of the carried "
            "reserves; the reviewed estimate replaces it in the close."]),
    ]
    for head, paras in sections:
        story.append(Paragraph(head, H2))
        for p in paras:
            story.append(Paragraph(p, CL))
    story.append(Spacer(1, 10))
    story.append(Paragraph("%s, FCAS, MAAA, Chief Actuary" % act, BODY))
    doc.build(story, onFirstPage=_frame(K.ORG, "Reserving basis 2026.1"),
              onLaterPages=_frame(K.ORG, "Reserving basis 2026.1"))


def write_platform_notes(path, names):
    who = K.PEOPLE["systems"]
    text = """# Claims platform extracts, bodily injury

Run by Claims Systems at 06:10 Central on 6 January 2027. Each extract holds the platform's records for bodily injury
(BI) claimant features as they stood at the run. Prepared by %s, Claims Systems.

Timestamps (fields ending `_at`) are ISO 8601 in UTC. `accident_date` and `reported_date` are calendar dates as the
handler recorded them.

## %s

One row per BI claimant feature held on the platform.

| Column | Meaning |
|---|---|
| feature_id | the platform's feature number |
| claim_number | the platform's claim (one accident; a claim can carry several claimant features) |
| legacy_ref | the legacy claim and feature number, for a feature that began in the legacy system |
| coverage | BI |
| state | the state of the policy |
| accident_date | the date of the accident |
| reported_date | the date the company first had notice of the claimant's injury |
| created_at | when the feature's platform record began: its creation, or its conversion or re-entry from the legacy system |
| status | OPEN, CLOSED or REOPENED, at the run |
| closed_at | the feature's latest closing at the run |
| reopened_at | the feature's latest reopening at the run |
| disposition | the reason recorded at the latest closing: SETTLED, DENIED, WITHDRAWN, NO_INJURY or DUPLICATE (a feature opened in error for a claimant who already has one; it is not a claimant feature) |

## %s

One row per event on a feature: CREATED, CONVERTED, CLOSED (with its disposition) or REOPENED, at `event_at`.

## %s

One row per payment the platform issued from 1 July 2024 to the run. `amount_cents` is in US cents. `method` is CHECK
or WIRE. A payment voided after issue has `status` VOIDED and its `voided_at`; a check issued to replace a voided one
names it in `reissue_of`.
""" % (who, names["features"], names["events"], names["pay"])
    open(path, "w", encoding="utf-8", newline="\n").write(text)


def write_conversion_notes(path):
    who = K.PEOPLE["systems"]
    text = """# Claims conversion, June 2024

Prepared by %s, Claims Systems, 8 July 2024.

- Bodily injury features open on the legacy claims system at the start of Saturday 29 June 2024 were converted to the
  claims platform over that weekend. The platform went live for all claims work on Monday 1 July 2024; no claims
  work was processed on 29 and 30 June.
- A converted feature keeps its legacy claim and feature number (the platform's `legacy_ref`), its accident date and
  its report date. Its platform record starts at the conversion, with a CONVERTED event.
- The legacy archive keeps every legacy feature. A converted feature's archive row shows disposition CV with the
  closing date 29 June 2024.
- Payments issued before 1 July 2024 stay in the legacy payment ledger under the legacy claim and feature number.
  From 1 July 2024 every payment is issued on the platform.
- The legacy system is read-only from 1 July 2024.
""" % who
    open(path, "w", encoding="utf-8", newline="\n").write(text)


def write_codes(path):
    text = ("LEGACY CLAIMS SYSTEM - BODILY INJURY FEATURE ARCHIVE\r\n"
            "DISPOSITION CODES (DISP)\r\n"
            "\r\n"
            "ST   SETTLED\r\n"
            "WD   WITHDRAWN BY CLAIMANT\r\n"
            "DN   DENIED\r\n"
            "NI   NO INJURY CLAIMED\r\n"
            "XD   OPENED IN ERROR - DUPLICATE OF ANOTHER FEATURE\r\n"
            "CV   CONVERTED TO CLAIMS PLATFORM 06/29/2024 - FEATURE CONTINUES ON PLATFORM\r\n"
            "\r\n"
            "CLS_DT    DATE OF THE FEATURE'S LAST CLOSING ON THE LEGACY SYSTEM\r\n"
            "REOPEN_CT TIMES THE FEATURE WAS REOPENED ON THE LEGACY SYSTEM\r\n"
            "DATES ARE MM/DD/YYYY. AMOUNTS IN THE PAYMENT LEDGER ARE DOLLARS; A VOID IS A NEGATIVE AMOUNT ON THE VOID DATE\r\n"
            "AGAINST THE CHECK IT CANCELS.\r\n")
    open(path, "w", encoding="ascii", newline="").write(text)


def write_ledger_notes(path, names):
    who = K.PEOPLE["controller"]
    text = """# General ledger extract: account 5311, 2026

Account 5311 is losses paid, auto bodily injury. The extract (%s) was run by Finance on 6 January 2027 for every
journal line posted from 1 January to 31 December 2026. Prepared by %s, Controller.

| Column | Meaning |
|---|---|
| journal_id, line | the journal and its line |
| posting_date | the date the line was posted |
| period | the accounting period, the posting date's month |
| source | CLAIMS FEED, TREASURY or MANUAL |
| batch_ref | the feed batch or the journal reference |
| state | the state segment |
| description | the journal's description |
| amount | US dollars; a reduction is negative |

The claims feed runs at 01:30 Central each night and posts the checks the claims platform issued the previous day,
net of the checks voided that day, one line per state. TREASURY lines are journals posted by Treasury; MANUAL lines
are journals posted by Finance.
""" % (names["gl"], who)
    open(path, "w", encoding="utf-8", newline="\n").write(text)


def write_calendar(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Close calendar 2026"
    ws.append(["Finance close calendar 2026"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(["Month", "Business days", "Last business day", "Ledger close (business day 5)", "Office holidays"])
    for c in ws[3]:
        c.font = Font(bold=True)
    for m in range(1, 13):
        d = dt.date(2026, m, 1)
        days = []
        while d.month == m:
            if W.business(d):
                days.append(d)
            d += dt.timedelta(days=1)
        ny, nm = (2027, 1) if m == 12 else (2026, m + 1)
        hol = sorted(h for h in W._HOL if h.year == 2026 and h.month == m)
        ws.append([dt.date(2026, m, 1).strftime("%B %Y"), len(days), days[-1], W.nth_business_day(ny, nm, 5),
                   ", ".join(h.strftime("%d %b") for h in hol) or None])
    for row in ws.iter_rows(min_row=4, max_row=15):
        row[2].number_format = "yyyy-mm-dd"
        row[3].number_format = "yyyy-mm-dd"
    for col, w in zip("ABCDE", (16, 14, 18, 28, 24)):
        ws.column_dimensions[col].width = w
    wb.save(path)


REGISTER = [
    ("Private passenger auto liability", "Bodily injury", "Loss", None),
    ("Private passenger auto liability", "Bodily injury", "Defence and cost containment", 52400),
    ("Private passenger auto liability", "Bodily injury", "Adjusting and other", 11900),
    ("Private passenger auto liability", "Property damage", "Loss", 24800),
    ("Private passenger auto liability", "Property damage", "Defence and cost containment", 1900),
    ("Private passenger auto liability", "Medical payments", "Loss", 6300),
    ("Auto physical damage", "Collision and comprehensive", "Loss", 9100),
    ("Homeowners", "All perils", "Loss", 41700),
    ("Homeowners", "All perils", "Adjusting expense", 6200),
]


def write_register(path, booked):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Preliminary YE2026"
    ws.append(["%s: preliminary year-end 2026 loss and adjusting expense reserves" % K.ORG])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Finance roll-forward for the close; replaced by the reviewed estimates. Amounts in $000, gross of "
               "reinsurance. Prepared 5 January 2027."])
    ws.append([])
    ws.append(["Line", "Coverage", "Component", "Preliminary 31 Dec 2026 ($000)"])
    for c in ws[4]:
        c.font = Font(bold=True)
    for line, cov, comp, amt in REGISTER:
        ws.append([line, cov, comp, int(round(booked / 1000.0)) if amt is None else amt])
    for row in ws.iter_rows(min_row=5, max_row=4 + len(REGISTER)):
        row[3].number_format = "#,##0"
    for col, w in zip("ABCD", (34, 30, 30, 30)):
        ws.column_dimensions[col].width = w
    wb.save(path)
