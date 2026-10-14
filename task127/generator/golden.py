"""The golden deliverables: the split workbook (the split of the 1,200 packages with the bed-days each hospital's
packages free, and the brokerage's referrals with care started within seven days) and the chart of each
hospital's occupied bed-days lost to delayed discharge last winter."""
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import core
import postprocess as PP

XLSX = "d2a_package_split_winter_2026-27.xlsx"
SVG = "delayed_bed_days_winter_2025-26.svg"
TRUST_SHORT = {"RMC": "Marchshire Hospitals", "RSM": "South Marches Healthcare"}


def whole(x):
    return int(x + 0.5)


def _split_sheet(ws, F):
    M = F["M"]
    bold = Font(bold=True)
    ws["A1"] = "Winter Discharge Fund 2026/27: split of the 1,200 discharge-to-assess packages"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Packages run 1 December 2026 to 28 February 2027. Sized on the patients the fund would have covered "
                "last winter (medically optimised 1 December 2025 to 28 February 2026), as allocation note SF-2026-14 "
                "sets out.")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:F2")
    ws.row_dimensions[2].height = 30
    head = ["Hospital", "Trust", "Covered patients, winter 2025/26", "Nights a package frees per covered patient",
            "Packages", "Occupied bed-days freed"]
    ws.append([])
    ws.append(head)
    fill = PatternFill("solid", fgColor="DCE6F1")
    for c in ws[4]:
        c.font = bold
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[4].height = 44
    for s in M["order"]:
        h = core.HOSPITALS[s]
        ws.append([h["name"], TRUST_SHORT[h["trust"]], M["caps"][s], float(M["val"][s]), M["packages"][s],
                   whole(M["bed"][s])])
    n = 4 + len(M["order"])
    ws.append(["Total", "", sum(M["caps"].values()), None, sum(M["packages"].values()), whole(M["total"])])
    top = Border(top=Side(style="thin"))
    for c in ws[n + 1]:
        c.font = bold
        c.border = top
    for row in ws.iter_rows(min_row=5, max_row=n + 1):
        row[2].number_format = "#,##0"
        row[3].number_format = "0.00"
        row[4].number_format = "#,##0"
        row[5].number_format = "#,##0"
    none = [core.HOSPITALS[s]["name"] for s in M["none"]]
    ws.append([])
    ws.append(["Hospitals with no packages: " + (", ".join(none) if none else "none") + "."])
    ws.cell(row=n + 3, column=1).font = bold
    s_low = M["none"][0] if M["none"] else M["order"][-1]
    ws.append(["%s's covered patients would have left %.2f nights sooner each with a package, the fewest of the five: "
               "most of them were also waiting for community equipment, and a package does not bring that forward."
               % (core.HOSPITALS[s_low]["name"], float(M["val"][s_low]))])
    ws.cell(row=n + 4, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=n + 4, start_column=1, end_row=n + 4, end_column=6)
    ws.row_dimensions[n + 4].height = 30
    small = Font(italic=True, size=9, color="555555")
    ws.append([])
    ws.append(["Covered patients exclude patients registered with Severnside practices and patients whose existing "
               "package restarted on discharge (allocation note SF-2026-14, section 3)."])
    ws.append(["Source: ICB data warehouse acute discharges (latest submissions) and ICES order lines, extracts of "
               "5 October 2026; GP practice list, September 2026."])
    for r in (n + 6, n + 7):
        ws.cell(row=r, column=1).font = small
        ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        ws.row_dimensions[r].height = 26
    for col, w in zip("ABCDEF", [26, 24, 15, 17, 11, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    ws.print_area = "A1:F%d" % (n + 7)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1


def _brokerage_sheet(ws, F):
    ws["A1"] = "Brokerage referrals for home care, winter 2025/26: care started within seven days of referral"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Referrals the council's brokerage received in each month, by the district of the address where care "
                "was to be delivered, and how many had care start no more than seven days after the day of referral "
                "(framework starts from the brokerage, spot purchases from the invoiced service start).")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:G2")
    ws.row_dimensions[2].height = 44
    ws.append([])
    months = core.ASK_MONTHS
    head = ["District"] + ["%s: care started within 7 days" % core.MONTH_LABEL[m] for m in months] + \
           ["%s: referrals received" % core.MONTH_LABEL[m] for m in months]
    ws.append(head)
    fill = PatternFill("solid", fgColor="DCE6F1")
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[4].height = 44
    for dist in core.DISTRICTS:
        ws.append([dist] + [F["B"][(dist, m)] for m in months] + [F["B_received"][(dist, m)] for m in months])
    for row in ws.iter_rows(min_row=5, max_row=7):
        for c in row[1:]:
            c.number_format = "#,##0"
    for col, w in zip("ABCDEFG", [16, 15, 15, 15, 14, 14, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B5"
    ws.append([])
    ws.append(["Source: Marchshire County Council brokerage referrals, spot-care invoice lines (finance) and postcode "
               "lookup, extracts of 5 October 2026."])
    ws.cell(row=ws.max_row, column=1).font = Font(italic=True, size=9, color="555555")
    ws.print_area = "A1:G%d" % ws.max_row
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1


def _notes_sheet(ws):
    lines = [
        ("Notes", True),
        ("Covered patients: aged 65 and over, discharged on pathway 1 from one of the five hospitals, needing a new "
         "package of home care (not a restart) and registered with a Marchshire practice (practices closed by merger "
         "counted with their successor), medically optimised 1 December 2025 to 28 February 2026.", False),
        ("Nights a package frees: for each covered patient, the nights from the later of the discharge ready date and "
         "the delivery of the equipment ordered for their discharge, to the day they left; averaged over the "
         "hospital's covered patients, patients a package would not have helped counted at nought.", False),
        ("Split: hospitals filled in order of nights freed per package, each up to its covered patients, until the "
         "1,200 packages are used. Bed-days freed are packages times nights per package, rounded to whole bed-days.",
         False),
        ("Spells are counted once, at their latest submission. Sources: ICB data warehouse acute discharges and ward "
         "stays and the ICES order lines (extracts of 5 October 2026); Marchshire County Council brokerage "
         "referrals, spot-care invoice lines and postcode lookup.", False),
        ("Rounding: nights a package frees are shown to two decimals and carried unrounded; packages, bed-days and "
         "referrals are whole numbers, each hospital's bed-days rounded once and the total taken from the unrounded "
         "figures.", False),
        ("Data as at the extracts of 5 October 2026. Prepared by System Flow, NHS Marchshire ICB, 12 October 2026.",
         False),
    ]
    for text, b in lines:
        ws.append([text])
        c = ws.cell(row=ws.max_row, column=1)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if b:
            c.font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 120
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 32


def write_xlsx(path, F):
    wb = Workbook()
    ws = wb.active
    ws.title = "Split"
    _split_sheet(ws, F)
    _brokerage_sheet(wb.create_sheet("Brokerage starts"), F)
    _notes_sheet(wb.create_sheet("Notes"))
    wb.save(path)
    PP.normalize_xlsx(path, "System Flow, NHS Marchshire ICB", "2026-10-12T09:27:13Z", "2026-10-12T15:11:52Z")


def write_svg(path, F):
    plt.rcParams["svg.hashsalt"] = "task127"
    plt.rcParams["font.family"] = "DejaVu Sans"
    A, AP = F["A"], F["AP"]
    order = sorted(core.SITES, key=lambda s: sum(A[(s, m)] for m in core.ASK_MONTHS))     # most at the top
    months = core.ASK_MONTHS
    cols = {(2025, 12): "#9ECAE1", (2026, 1): "#3182BD", (2026, 2): "#08306B"}
    fig, ax = plt.subplots(figsize=(10.0, 6.4), dpi=100)
    h = 0.26
    for j, m in enumerate(months):
        ys = [i - (j - 1) * h for i in range(len(order))]
        vals = [A[(s, m)] for s in order]
        bars = ax.barh(ys, vals, height=h * 0.9, color=cols[m], label=core.MONTH_LABEL[m])
        for b, s, v in zip(bars, order, vals):
            ax.text(v + 18, b.get_y() + b.get_height() / 2, "{:,} ({:,})".format(v, AP[(s, m)]), va="center",
                    ha="left", fontsize=8, color="#222222")
    ax.set_yticks(range(len(order)))
    ax.set_yticklabels([core.HOSPITALS[s]["name"] for s in order], fontsize=9.5)
    ax.set_xlabel("Occupied bed-days lost to delayed discharge")
    ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, p: "{:,.0f}".format(v)))
    ax.set_xlim(0, max(A.values()) * 1.18)
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    ax.tick_params(axis="y", length=0)
    ax.legend(frameon=False, fontsize=8.5, loc="lower right")
    fig.text(0.012, 0.955, "Bed-days lost to delayed discharge by acute site, December 2025 to February 2026",
             fontsize=12, fontweight="bold", color="#111111")
    fig.text(0.012, 0.918, "Patients aged 65 and over, every discharge pathway. Each bar is a month at a site; "
             "labels give bed-days and, in brackets, the patients delayed.", fontsize=8.6, color="#333333")
    fig.text(0.012, 0.045, "A bed-day is a night spent in hospital after the discharge ready date, counted at the site "
             "whose ward the patient was on at midnight and dated to the evening the night began.", fontsize=7.2,
             color="#555555")
    fig.text(0.012, 0.018, "Source: ICB data warehouse, acute discharges (latest submissions) and ward stays, "
             "extracts of 5 October 2026. System Flow, NHS Marchshire ICB.", fontsize=7.2, color="#555555")
    fig.subplots_adjust(left=0.19, right=0.98, top=0.87, bottom=0.16)
    fig.savefig(path, format="svg", metadata={"Date": None, "Creator": "System Flow, NHS Marchshire ICB"})
    plt.close(fig)
    svg = open(path, encoding="utf-8").read().replace('id="matplotlib.axis_', 'id="axis_')
    assert "matplotlib" not in svg.lower()
    open(path, "w", encoding="utf-8").write(svg)


def write(golden_dir, F):
    os.makedirs(golden_dir, exist_ok=True)
    write_xlsx(os.path.join(golden_dir, XLSX), F)
    write_svg(os.path.join(golden_dir, SVG), F)
