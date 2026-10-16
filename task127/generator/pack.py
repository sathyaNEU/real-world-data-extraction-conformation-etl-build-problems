"""Writers for target/: the ICB warehouse's acute discharges and the trusts' ward stays with the ward reference, the
GP practice list, the community equipment service's order lines, the council's brokerage referrals, spot-care
invoice lines and postcode lookup, the fund guidance, the extract notes and the data request log, and the
distractors (the national weekly return, the daily bed census, the reablement outcomes)."""
import csv
import datetime as dt
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import core
import docs as D
import postprocess as PP
import records as RC
import world as W

FILES = {
    "spells": "acute_discharges_65plus_2024-12_2026-09.csv",
    "wards": "acute_ward_stays_2024-12_2026-09.csv",
    "wardref": "ward_reference.csv",
    "acute_notes": "acute_extract_notes.txt",
    "practices": "gp_practices_2026-09.csv",
    "ices": "ices_order_lines_2024-12_2026-09.csv",
    "ices_notes": "ices_export_notes.txt",
    "brokerage": "brokerage_referrals_2024-12_2026-09.csv",
    "spot": "spot_care_invoice_lines_2025-26.csv",
    "postcodes": "postcode_districts_marchshire.csv",
    "council_notes": "council_extract_notes.txt",
    "guidance": "winter_discharge_fund_2026-27.pdf",
    "weekly": "discharge_delays_weekly_return_2024-12_2026-09.csv",
    "census": "bed_census_daily_2024-12_2026-09.xlsx",
    "reablement": "reablement_91_day_outcomes_2024-26.xlsx",
    "log": "data_request_log.csv",
}
DISTRACTORS = ["weekly", "census", "reablement"]
TRUST_NAME = core.TRUSTS
SITE_NAME = {s: h["name"] for s, h in core.HOSPITALS.items()}
DT = "%Y-%m-%d %H:%M"


def _w(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)
    return len(rows)


def spaced(nhs):
    return "%s %s %s" % (nhs[:3], nhs[3:6], nhs[6:])


def practice_at(p, admitted):
    """The practice the patient was registered with at admission: the trusts' systems take it from the national
    register when the patient is admitted, so a practice closed by merger appears only on spells admitted before the
    merger and the successor's code on every later one."""
    code = p["practice"]
    if code in W.MERGED and admitted >= W.MERGED[code][2]:
        return W.MERGED[code][1]
    return code


def write_spells(R, subs, target):
    hn = RC.hospital_numbers(R["spells"])
    rows = []
    for s in sorted(subs, key=lambda x: x["sub_id"]):
        p = R["persons"][s["nhs"]]
        r = W.rng_for("adm", s["spell"])
        method = "Elective" if r.random() < 0.07 else "Emergency"
        rows.append([s["sub_id"], s["sub_d"].isoformat(), s["trust"], s["site"], s["spell"], s["nhs"],
                     hn[(s["trust"], s["nhs"])], s["age"], practice_at(p, s["adm_t"].date()), p["postcode"].split()[0],
                     method,
                     s["adm_t"].strftime(DT), s["ready"].isoformat(), s.get("dis_sub", s["dis_t"]).strftime(DT),
                     RC.DEST[s["pw_sub"]],
                     s["pw_sub"], RC.P1_CARE[s["p1"]] if s["pw_sub"] == 1 else ""])
    return _w(os.path.join(target, FILES["spells"]),
              ["submission_id", "submitted_on", "provider_code", "site_code", "spell_number", "nhs_number",
               "local_patient_id", "age_on_admission", "gp_practice_code", "postcode_district", "admission_method",
               "admission_datetime", "discharge_ready_date", "discharge_datetime", "discharge_destination",
               "discharge_pathway", "p1_care_package"], rows)


def write_wards(wstays, target):
    # South Marches' bed management system writes day/month/year
    f = lambda w, x: x.strftime("%d/%m/%Y %H:%M" if w["trust"] == "RSM" else DT)
    rows = [[w["trust"], w["spell"], w["ward"], f(w, w["start"]), f(w, w["end"])]
            for w in sorted(wstays, key=lambda x: (x["trust"], x["spell"], x["start"]))]
    return _w(os.path.join(target, FILES["wards"]),
              ["provider_code", "spell_number", "ward_code", "ward_start", "ward_end"], rows)


def write_wardref(target):
    rows = []
    for site, ws in RC.WARDS.items():
        tr = core.HOSPITALS[site]["trust"]
        for code, name, spec, beds in ws:
            rows.append([code, name, spec, site, SITE_NAME[site], tr, TRUST_NAME[tr], beds])
    return _w(os.path.join(target, FILES["wardref"]),
              ["ward_code", "ward_name", "specialty", "site_code", "site_name", "provider_code", "provider_name",
               "beds"], rows)


def write_practices(target):
    rows = []
    for code, (name, icb, area) in sorted(W.PRACTICES.items()):
        rows.append([code, name, W.ICB_NAME[icb], area, "Active", "", ""])
    for code, (name, succ, when) in sorted(W.MERGED.items()):
        rows.append([code, name, W.ICB_NAME["Marchshire"], W.PRACTICES[succ][2], "Closed (merged)", succ,
                     when.isoformat()])
    rows.sort()
    return _w(os.path.join(target, FILES["practices"]),
              ["practice_code", "practice_name", "icb", "locality", "status", "successor_code", "status_date"], rows)


def write_ices(eq, target):
    rows, line = [], {}
    for x in eq:
        line[x["order_no"]] = line.get(x["order_no"], 0) + 1
        f = lambda d: d.isoformat() if d else ""
        rows.append([x["order_no"], line[x["order_no"]], spaced(x["nhs"]), x["postcode"], x["item"], x["route"],
                     f(x["ordered"]), f(x["required"]), f(x["delivered"]), f(x["collected"]), x["depot"], x["status"]])
    return _w(os.path.join(target, FILES["ices"]),
              ["order_no", "line_no", "nhs_number", "delivery_postcode", "item", "route", "ordered_date",
               "required_by", "delivered_date", "collected_date", "depot", "status"], rows)


def styled(pc, style):
    if style == 1:
        return pc.replace(" ", "")
    if style == 2:
        return pc.lower()
    if style == 3:
        return pc + " "
    return pc


def write_brokerage(refs, target):
    rows = []
    for x in refs:
        p = x["person"]
        rows.append([x["ref_no"], x["ref_t"].strftime(DT), x["source"], SITE_NAME[x["site"]] if x["site"] else "",
                     p["person_id"], spaced(p["nhs"]) if x["nhs_shown"] else "", styled(p["postcode"], x["pc_style"]),
                     x["status"], x["provider"], x["start"].isoformat() if x["start"] else ""])
    return _w(os.path.join(target, FILES["brokerage"]),
              ["referral_no", "received_at", "source", "referring_site", "person_id", "nhs_number", "care_postcode",
               "status", "provider", "start_date"], rows)


def write_spot(lines, target):
    uk = lambda d: d.strftime("%d/%m/%Y")
    rows = [[x["invoice"], x["provider"], "BRK%07d" % int(x["ref_no"][2:]), x["person"]["person_id"], uk(x["start"]),
             uk(x["week_end"]),
             "%.1f" % x["hours"], "%.2f" % x["rate"], "%.2f" % round(x["hours"] * x["rate"] + 1e-9, 2)] for x in lines]
    return _w(os.path.join(target, FILES["spot"]),
              ["invoice_no", "supplier", "brokerage_ref", "client_id", "service_start", "week_ending", "hours", "rate",
               "amount"], rows)


def write_postcodes(target):
    rows = []
    lookup = W.lookup_postcodes()
    for dist in core.DISTRICTS:
        for pc in lookup[dist]:
            rows.append([pc, W.WARDS_ELECT[pc.split()[0]], dist])
    rows.sort()
    return _w(os.path.join(target, FILES["postcodes"]), ["postcode", "electoral_ward", "district"], rows)


def write_weekly(spells, target):
    wr = RC.weekly_return(spells)
    rows = []
    for (tr, we), v in sorted(wr.items(), key=lambda x: (x[0][1], x[0][0])):
        if not dt.date(2024, 12, 8) <= we <= dt.date(2026, 9, 27):
            continue
        rows.append([we.isoformat(), tr, TRUST_NAME[tr]] + v)
    head = ["week_ending", "provider_code", "provider_name"]
    for pw in range(4):
        head += ["p%d_discharged" % pw, "p%d_nights_after_ready" % pw]
    return _w(os.path.join(target, FILES["weekly"]), head, rows)


def _sheet(ws, head, rows, widths, fmt=None):
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCE6F1")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for col, f in (fmt or {}).items():
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for c in row:
                c.number_format = f
    ws.freeze_panes = "A2"


def write_census(wstays, target):
    occ = RC.occupancy(wstays)
    wb = Workbook()
    ws = wb.active
    ws.title = "Midnight census"
    rows = [[d, s, SITE_NAME[s], o, e, n] for d, s, o, e, n in occ]
    _sheet(ws, ["census_date", "site_code", "site_name", "beds_open", "escalation_beds_open", "occupied_at_midnight"],
           rows, [12, 10, 26, 11, 13, 13], {1: "dd/mm/yyyy"})
    ns = wb.create_sheet("Notes")
    for line in ["General and acute beds, every age, counted at midnight; the census date is the date the night begins.",
                 "beds_open includes escalation beds opened on the night.",
                 "Source: the two trusts' daily bed state returns to the system coordination centre."]:
        ns.append([line])
    ns.column_dimensions["A"].width = 110
    path = os.path.join(target, FILES["census"])
    wb.save(path)
    PP.normalize_xlsx(path, "System Coordination Centre", "2026-10-02T07:52:18Z", "2026-10-02T08:16:41Z")
    return len(rows)


def write_reablement(target):
    rr = RC.reablement()
    wb = Workbook()
    ws = wb.active
    ws.title = "Outcomes"
    qname = {1: "Jan to Mar", 2: "Apr to Jun", 3: "Jul to Sep", 4: "Oct to Dec"}
    rows = [["%s %d" % (qname[x["quarter"]], x["year"]), x["district"], x["discharged"], x["at_home_91"],
             round(x["at_home_91"] / x["discharged"], 3)] for x in rr]
    _sheet(ws, ["quarter_of_discharge", "district", "discharged_into_reablement", "at_home_91_days_later",
                "share_at_home"], rows, [12, 14, 16, 16, 12], {5: "0.0%"})
    ns = wb.create_sheet("Notes")
    for line in ["Older people (65 and over) discharged from hospital into the council's reablement service, and those "
                 "still at home 91 days after discharge, by district and calendar quarter of discharge.",
                 "Quarters are calendar quarters (Q1 January to March). Source: Marchshire County Council, adult social "
                 "care outcomes return."]:
        ns.append([line])
    ns.column_dimensions["A"].width = 110
    path = os.path.join(target, FILES["reablement"])
    wb.save(path)
    PP.normalize_xlsx(path, "Adult Social Care Performance", "2026-08-13T16:22:07Z", "2026-08-14T10:04:39Z",
                      company=core.COUNCIL)
    return len(rows)


def write_docs(target, counts):
    D.guidance_pdf(os.path.join(target, FILES["guidance"]))
    for k, fn in (("acute_notes", D.acute_notes), ("ices_notes", D.ices_notes), ("council_notes", D.council_notes)):
        open(os.path.join(target, FILES[k]), "w", encoding="utf-8").write(fn())
    _w(os.path.join(target, FILES["log"]), D.LOG_HEAD, D.log_rows(FILES, counts))
