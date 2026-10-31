"""task451 checks: run on the built pack (target/) and goldens. Every assertion prints; any failure fails the build.
Writes nothing but its report line; golden.py writes figures.json."""
import datetime as dt
import json
import os
import re
import sys

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import calc as C        # noqa: E402
import golden as G      # noqa: E402
import pack as P        # noqa: E402
import world as W       # noqa: E402

TASK = os.path.dirname(HERE)
COUNT, FAILS = [0], []

GOLDEN_SET = ["SIM", "KRH", "HPM", "BCN"]
EXPECT = {
    "G1 received, all development": GOLDEN_SET,
    "G2 received, latest 12 months": GOLDEN_SET,
    "G3 paid, after the go-live month": GOLDEN_SET,
    "G4 received, by claim type": GOLDEN_SET,
    "G5 paid (latest versions), after the go-live month": GOLDEN_SET,
    "R1 paid as extracted": ["SIM", "BCN"],
    "paid, latest 4 months": GOLDEN_SET + ["ALM"],
    "paid, latest 6 months": GOLDEN_SET + ["ALM"],
    "paid, latest 12 months": GOLDEN_SET + ["ALM", "CFP", "RPG", "THA", "JHP"],
    "R0 paid, all development": GOLDEN_SET + ["ALM", "CFP", "RPG", "THA", "JHP"],
}
GOLD_MARGIN = 0.010      # every golden route at least 1 per cent from the corridor line for every partner
RUNG_MARGIN = 0.005      # every other route's landing at least 0.5 per cent from it

# quoted sentences: each SF row's ruling sentence, present once in its document
QUOTES = [
    ("SF2", "agreement", "exceeds its Target PMPM in Exhibit 1 by more than 2.0 per cent"),
    ("SF3", "agreement", "with a first date of service in the Performance Year"),
    ("SF4", "agreement", "paid through the last business day of January following the year"),
    ("SF5", "agreement", "the Fund's paid amount for covered medical services"),
    ("SF6", "agreement", "eligible for the Fund's medical benefits and appears on the Fund's attribution roster"),
    ("SF7", "agreement", "the amount above $200,000 is excluded"),
    ("SF8", "agreement", "Outpatient prescription drugs dispensed by retail, mail-order and specialty pharmacies are excluded"),
    ("SF9", "notice", "so a claim's value is the paid amount of its latest version"),
    ("SF10", "agreement", "attributed to the Partner in the month of that first date of service"),
    ("SF11", "agreement", "are made on unrounded figures"),
    ("SF12", "agreement", "whenever the claims are paid"),
    ("SF13", "agreement", "attributed to the Partner in the month of that first date of service"),
    ("SF14", "agreement", "shared in proportion to the Member's amounts with each"),
]
BANNED = re.compile(r"\blag\b|lagged|triangle|completion|\bIBNR\b|incurred but not reported|chain[- ]ladder|run-?out|"
                    r"develop(ment|ed)? factor|\breserve\b|turnaround|days to pay|faster|slower|speed", re.I)


def ok(cond, msg):
    COUNT[0] += 1
    if cond:
        print("ok    " + msg)
    else:
        FAILS.append(msg)
        print("FAIL  " + msg)


def text_of(path):
    if path.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join(p.extract_text() for p in PdfReader(path).pages)
    if path.endswith(".docx"):
        import docx
        return "\n".join(p.text for p in docx.Document(path).paragraphs)
    if path.endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        return "\n".join(" ".join(str(c) for c in r if c is not None) for ws in wb for r in ws.iter_rows(values_only=True))
    return open(path, encoding="utf-8", errors="replace").read()


def x_of(pm, tg):
    return {c: pm[c] / ((1 + W.CORRIDOR) * tg[c]) for c in W.CODES}


def flags_of(pm, tg):
    return sorted(c for c in W.CODES if pm[c] > tg[c] * (1 + W.CORRIDOR))


def run(target, golden, world=None, meta_path=None):
    names = sorted(os.listdir(target))
    texts = {n: text_of(os.path.join(target, n)) for n in names}
    meta = json.load(open(meta_path or os.path.join(TASK, "metadata.json")))
    X = G.compute(target)
    L, tg = X["L"], X["targets"]
    # ------------------------------------------------------------------ input gates
    ok(len(names) >= 10, "ten or more files (%d)" % len(names))
    fmts = {n.rsplit(".", 1)[1] for n in names}
    ok(len(fmts) >= 3, "three or more formats (%s)" % sorted(fmts))
    ok(len(L["claims"]) >= 25000, "the spine has 25,000 rows or more (%d)" % len(L["claims"]))
    ok(len(meta["distractor_files"]) >= 2 and all(d in names for d in meta["distractor_files"]), "two or more distractors named in metadata.json")
    ok(not any("distractor" in n.lower() for n in names), "no file name says distractor")
    ok(sorted(names) == sorted(P.FILES.values()), "the pack is exactly the files the generator names")
    describing = re.compile(r"dictionar|manifest|readme|read_me|field_?notes|field_?guide|codebook|schema|definition|"
                            r"extract_?notes|export_?notes|data_?notes|layout|release_?notes", re.I)
    ok(not [n for n in names if describing.search(n)], "section 3: no dictionary, file layout, readme or data notes ships")
    roster_codes = set(L["roster"]["partner_code"])
    ok(sorted(tg) == sorted(W.CODES) and roster_codes <= set(tg),
       "SF1 the accountable care partners are the groups of Exhibit 1 (%d), and the roster names no other" % len(tg))

    # ------------------------------------------------------------------ the call and every route
    for nm, want in EXPECT.items():
        got = X["flags"][nm]
        ok(sorted(got) == sorted(want), "%s names %s (%s)" % (nm, ",".join(sorted(want)), ",".join(got)))
        xs = x_of(X["pmpm"][nm], tg)
        m = min(abs(v - 1) for v in xs.values())
        need = GOLD_MARGIN if nm in G.FAMILY else RUNG_MARGIN
        ok(m >= need, "%s: every partner at least %.1f%% from the corridor line (closest %.2f%%)" % (nm, 100 * need, 100 * m))
    ok(len(set(tuple(X["flags"][nm]) for nm in G.FAMILY)) == 1, "every golden route names the same partners")
    ok(len({tuple(sorted(v)) for v in EXPECT.values()}) == 4, "four distinct sets across the routes (golden, R1, the short paid ladders, the long ones)")
    ok(abs(sum(X["pmpm"]["R0 paid, all development"][c] * X["mm"][c] for c in W.CODES) /
           sum(X["pmpm"][G.GOLD][c] * X["mm"][c] for c in W.CODES) - 1) > 0.06, "the default overstates the book by more than 6 per cent")

    # ------------------------------------------------------------------ the reflex sweep: every paid-date habit lands wrong
    tri, amt, amt_t, exc0, mm0 = X["tri"], X["amt"], X["amt_t"], X["excess"], X["mm"]

    def simple(T, keep=None):
        n = T.shape[0]
        f = np.ones(n)
        for L_ in range(n - 1):
            r = [T[i, L_ + 1] / T[i, L_] for i in range(n - L_ - 1) if T[i, L_] > 0 and (keep is None or keep(C.S0 + i + L_ + 1))]
            f[L_] = float(np.mean(r)) if r else 1.0
        return f

    def named(f):
        return flags_of(C.tcoc(f, amt, amt_t, exc0, mm0), tg)

    tt = C.by_type_triangles(L["claims"], X["fc"])
    wrong = {
        "paid, all development without the go-live month's payments": C.completion(C.factors(tri["paid"], lambda m: m != C.GOLIVEM)),
        "paid, latest 12 months without the go-live month's payments": C.completion(C.factors(tri["paid"], lambda m: m > C.CUTM - 12 and m != C.GOLIVEM)),
        "paid, latest 5 months": C.completion(C.factors(tri["paid"], lambda m: m > C.CUTM - 5)),
        "paid, payments before the go-live only": C.completion(C.factors(tri["paid"], lambda m: m < C.GOLIVEM)),
        "paid, simple average of link ratios": C.completion(simple(tri["paid"])),
        "paid, by claim type": {t: C.completion(C.factors(tt[t]["paid"])) for t in tt},
        "paid, by claim type, latest 12 months": {t: C.completion(C.factors(tt[t]["paid"], lambda m: m > C.CUTM - 12)) for t in tt},
        "paid on latest versions, all development": C.completion(C.factors(tri["paid_final"])),
        "paid on latest versions, latest 6 months": C.completion(C.factors(tri["paid_final"], lambda m: m > C.CUTM - 6)),
        "paid on latest versions, from the go-live month (its backlog included)": C.completion(C.factors(tri["paid_final"], lambda m: m >= C.GOLIVEM)),
    }
    for nm, f in wrong.items():
        got = named(f)
        ok(got != sorted(GOLDEN_SET) and len(got) >= 5, "reflex %s names %d partners, not the four" % (nm, len(got)))
    right = {
        "received, simple average of link ratios": C.completion(simple(tri["received"])),
        "received, latest 6 months": C.completion(C.factors(tri["received"], lambda m: m > C.CUTM - 6)),
        "paid, the latest 2 payment months": C.completion(C.factors(tri["paid"], lambda m: m > C.CUTM - 2)),
    }
    for nm, f in right.items():
        ok(named(f) == sorted(GOLDEN_SET), "%s names the four" % nm)

    # ------------------------------------------------------------------ the forcing evidence
    lg = X["lags"]
    ok(lg["r2p_median_before"] >= 15 and lg["r2p_median_after"] <= 3,
       "receipt to payment: median %.0f days before go-live, %.0f after" % (lg["r2p_median_before"], lg["r2p_median_after"]))
    ok(abs(lg["s2r_median_2025"] - lg["s2r_median_after"]) <= 2,
       "service to receipt unchanged: median %.0f days in 2025, %.0f since go-live" % (lg["s2r_median_2025"], lg["s2r_median_after"]))
    fr_all = C.factors(X["tri"]["received"])
    fr_old = C.factors(X["tri"]["received"], lambda m: m < C.GOLIVEM)
    fp_old = C.factors(X["tri"]["paid"], lambda m: m < C.GOLIVEM)
    fp_new = C.factors(X["tri"]["paid"], lambda m: m > C.GOLIVEM)
    ok(abs(fr_all[0] / fr_old[0] - 1) < 0.03 and abs(fr_all[1] / fr_old[1] - 1) < 0.02,
       "the received-date development holds across the go-live (first link %.3f against %.3f before it)" % (fr_all[0], fr_old[0]))
    ok(fp_old[0] / fp_new[0] > 1.25, "the paid-date development does not (first link %.3f before go-live, %.3f after)" % (fp_old[0], fp_new[0]))

    # ------------------------------------------------------------------ R2: every Layer 1 item neutral to the call
    cl, fc = L["claims"], X["fc"]
    att, elig = X["att"], X["elig"]

    def call_with(cl_=None, fc_=None, mm_=None, excess_=None):
        cl_ = cl if cl_ is None else cl_
        fc_ = fc if fc_ is None else fc_
        d26 = C.frame2026(fc_, att)
        exc, _ = C.truncation(d26)
        tri = C.triangles(cl_, fc_)
        f = C.completion(C.factors(tri["received"]))
        amt, amt_t = C.partner_amounts(d26)
        pm = C.tcoc(f, amt, amt_t, exc if excess_ is None else excess_, mm_ or X["mm"])
        return flags_of(pm, tg)

    summed = fc.copy()
    summed["paid_amount"] = cl.groupby("claim_number")["paid_amount"].sum().reindex(summed["claim_number"]).to_numpy()
    ok(call_with(fc_=summed) == sorted(GOLDEN_SET), "R2 every version's paid amount summed: same call")
    v1 = cl[cl["claim_version"] == 1].copy()
    v1f = fc.copy()
    v1f["paid_amount"] = v1.set_index("claim_number")["paid_amount"].reindex(v1f["claim_number"]).to_numpy()
    ok(call_with(fc_=v1f) == sorted(GOLDEN_SET), "R2 each claim's first version only: same call")
    mm_roster = C.member_months(L["roster"], elig, require_eligible=False)
    ok(call_with(mm_=mm_roster) == sorted(GOLDEN_SET), "R2 member months from the roster alone (retroactive terminations kept): same call")
    first_ver = L["elig"].sort_values(["span_id", "span_version"]).groupby("span_id").head(1)
    mm_v1 = C.member_months(L["roster"], C.eligible_months(first_ver), require_eligible=True)
    ok(call_with(mm_=mm_v1) == sorted(GOLDEN_SET), "R2 member months on each span's first version: same call")
    tot_m = X["d26"].groupby("member_id")["paid_amount"].sum()
    over_m = tot_m[tot_m > W.TRUNCATION]
    d_tr = X["d26"].copy()
    d_tr["paid_amount"] = d_tr["paid_amount"] * d_tr["member_id"].map(lambda m: W.TRUNCATION / over_m[m] if m in over_m.index else 1.0)
    a_tr, at_tr = C.partner_amounts(d_tr)
    for nm in G.FAMILY:
        pm_tr = C.tcoc(X["cf"][nm], a_tr, at_tr, {c: 0.0 for c in W.CODES}, X["mm"])
        ok(flags_of(pm_tr, tg) == sorted(GOLDEN_SET), "closure 13 %s with the $200,000 taken off each member's claims before completing: same call" % nm)
    ok(all(X["excess"][c] == 0 for c in GOLDEN_SET), "no owing partner carries a truncated member, so its share of target is the same in either order")

    # ------------------------------------------------------------------ R3: the three rules executed
    ok(X["flags"]["R0 paid, all development"] != X["call"], "R3a the default path (paid-date chain ladder, every filed rule applied) names %d partners, not the golden %d" % (len(X["flags"]["R0 paid, all development"]), len(X["call"])))
    ok(X["flags"]["G1 received, all development"] == sorted(GOLDEN_SET), "R3b the default path with only the development basis changed names the golden four")
    ok(X["flags"]["R1 paid as extracted"] != X["call"], "R3b the lower rung (paid as extracted) is still wrong")
    ok(not any(col in L["claims"].columns for col in ("platform", "system", "adjudication_system", "processing_system")),
       "no column marks the adjudication platform")

    # ------------------------------------------------------------------ semantic forks
    for sf, doc, q in QUOTES:
        t = texts[P.FILES[doc]].replace("\n", " ")
        ok(re.sub(r"\s+", " ", q) in re.sub(r"\s+", " ", t), "%s quoted sentence present in %s" % (sf, P.FILES[doc]))
    xs = x_of(X["pmpm"][G.GOLD], tg)
    ok(all(abs(v - 1) > 1e-4 for v in xs.values()), "SF2 more than against at or over: same call")
    ok(all(abs(X["pmpm"][G.GOLD][c] / tg[c] - 1.02) > 0.0005 for c in W.CODES), "SF11 rounding first to a cent or a tenth of a per cent: same call")

    # ------------------------------------------------------------------ the asks
    agr = re.sub(r"\s+", " ", texts[P.FILES["agreement"]].replace("\n", " "))
    for q in ("prescriptions filled in the Performance Year in the Members' Attributed Member Months with the Partner, reversals included whenever processed",
              "the Fund's paid amount for the professional services in its Total Cost of Care, before B.1.5"):
        ok(q in agr, "ask pin present in the agreement: %s" % q[:60])
    rx = X["rx"]
    rx_pc = C.ask_pharmacy(L, att, "person_code", elig)
    rx_ch = C.ask_pharmacy(L, att, "cardholder", elig)
    rx_ro = C.ask_pharmacy(L, att, "dob")
    for c in W.CODES:
        frac = rx[c] / 1000 - np.floor(rx[c] / 1000)
        ok(abs(frac - 0.5) * 1000 >= 150, "Ask A %s pharmacy %.2f at least $150 from a rounding edge" % (c, rx[c]))
        ok(G.k1(rx_pc[c]) != G.k1(rx[c]), "Ask A %s: the person-code join files %d against %d" % (c, G.k1(rx_pc[c]), G.k1(rx[c])))
        ok(G.k1(rx_ch[c]) != G.k1(rx[c]), "Ask A %s: the cardholder join files %d against %d" % (c, G.k1(rx_ch[c]), G.k1(rx[c])))
    ok(sum(G.k1(rx_ro[c]) != G.k1(rx[c]) for c in W.CODES) >= 6,
       "Ask A: the roster alone (fills in months the member is not eligible kept) misses %d of 12 partners" % sum(G.k1(rx_ro[c]) != G.k1(rx[c]) for c in W.CODES))
    sh = X["share"]
    sh_snap = C.ask_leakage(L, X["d26"], X["group_npi"], "snapshot")
    sh_since = C.ask_leakage(L, X["d26"], X["group_npi"], "since")
    for c in W.CODES:
        frac = sh[c] * 10 - np.floor(sh[c] * 10)
        ok(abs(frac - 0.5) / 10 >= 0.02, "Ask B %s share %.4f at least 0.02 points from a rounding edge" % (c, sh[c]))
        ok("%.1f" % sh_snap[c] != "%.1f" % sh[c], "Ask B %s: the January roster read for the year files %.1f against %.1f" % (c, sh_snap[c], sh[c]))
        ok("%.1f" % sh_since[c] != "%.1f" % sh[c], "Ask B %s: the roster with its since dates files %.1f against %.1f" % (c, sh_since[c], sh[c]))
    # the cracker (golden call, the natural joins) and the mirror (the default call, the same joins) miss every ask row
    crack = sum(G.k1(rx_pc[c]) == G.k1(rx[c]) for c in W.CODES) + sum("%.1f" % sh_since[c] == "%.1f" % sh[c] for c in W.CODES)
    ok(crack == 0, "cracker sheet: %d of 24 ask rows hit" % crack)

    # ------------------------------------------------------------------ separation: nothing the asks turn on moves the call
    scram = fc.copy()
    scram["rendering_npi"] = scram["rendering_npi"].sample(frac=1.0, random_state=7).to_numpy()
    scram["billing_npi"] = scram["billing_npi"].sample(frac=1.0, random_state=8).to_numpy()
    ok(call_with(fc_=scram) == sorted(GOLDEN_SET), "separation: scrambling every claim's NPIs leaves the call")
    ok(not any(k in texts[P.FILES["agreement"]] for k in ("person code", "group NPI on the claim", "billed under")),
       "no filed sentence names either ask's device")

    # ------------------------------------------------------------------ documents
    for n in names:
        if n.endswith(".csv"):
            continue
        hit = BANNED.findall(texts[n])
        ok(not hit, "%s names no development, lag or payment speed (%s)" % (n, hit[:3]))
    for p_ in W.PEOPLE.values():
        ok(any(p_[0] in texts[n] for n in names), "persona %s appears in the pack" % p_[0])
    ok(len(set(tg.values())) == len(tg) and all(abs(v - round(v, 0)) >= 0.05 for v in tg.values()), "no target on a whole dollar, none repeated")
    ok(L["claims"]["claim_number"].nunique() == (L["claims"]["claim_version"] == 1).sum(), "one first version per claim")
    ok(not L["claims"].duplicated(["claim_number", "claim_version"]).any(), "unique claim and version keys")
    gd = sorted(os.listdir(golden))
    ok(gd == sorted([G.XLSX, G.PNG]), "the goldens are the two files the prompt names (%s)" % gd)
    print("checks: %d assertions, %d failed" % (COUNT[0], len(FAILS)))
    return not FAILS


if __name__ == "__main__":
    tgt = sys.argv[1] if len(sys.argv) > 1 else os.path.join(TASK, "target")
    sys.exit(0 if run(tgt, os.path.join(TASK, "golden")) else 1)
