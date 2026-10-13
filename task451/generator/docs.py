"""The pack's documents: the partner agreement's Schedule B with Exhibit 1 (each partner's 2026 target), Keswold's
extract dictionary, the Fund's data notes, Corvane's claims file layout and the Fund's stop-loss policy summary."""
import datetime as dt

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import LETTER                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import inch                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import world as W   # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("LiberationSerif", "Regular"), ("LiberationSerif-Bold", "Bold"), ("LiberationSans", "Regular"),
              ("LiberationSans-Bold", "Bold")):
    pdfmetrics.registerFont(TTFont(nm, FONT_DIR + nm.split("-")[0] + "-" + f + ".ttf"))
INK = colors.HexColor("#1d1f21")
GREY = colors.HexColor("#5a5f66")

BODY = ParagraphStyle("body", fontName="LiberationSerif", fontSize=10.5, leading=14.5, textColor=INK, spaceAfter=7)
NUM = ParagraphStyle("num", parent=BODY, leftIndent=28, firstLineIndent=-28)
SUBNUM = ParagraphStyle("subnum", parent=BODY, leftIndent=46, firstLineIndent=-18)
H1 = ParagraphStyle("h1", fontName="LiberationSerif-Bold", fontSize=13.5, leading=17, textColor=INK, spaceAfter=4,
                    alignment=1)
SUB = ParagraphStyle("sub", fontName="LiberationSerif", fontSize=10, leading=13, textColor=GREY, spaceAfter=14, alignment=1)
H2 = ParagraphStyle("h2", fontName="LiberationSerif-Bold", fontSize=11, leading=15, textColor=INK, spaceBefore=8, spaceAfter=4)
CELL = ParagraphStyle("cell", fontName="LiberationSans", fontSize=8.8, leading=11, textColor=INK)
CELLB = ParagraphStyle("cellb", parent=CELL, fontName="LiberationSans-Bold")
SMALL = ParagraphStyle("small", fontName="LiberationSans", fontSize=8.5, leading=11, textColor=GREY)


def _doc(path, title, author, footer):
    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFont("LiberationSans", 7.5)
        canvas.setFillColor(GREY)
        canvas.drawString(1.0 * inch, 0.6 * inch, footer)
        canvas.drawRightString(LETTER[0] - 1.0 * inch, 0.6 * inch, "Page %d" % doc.page)
        canvas.restoreState()
    d = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=1.0 * inch, rightMargin=1.0 * inch, topMargin=0.9 * inch,
                          bottomMargin=0.95 * inch, title=title, author=author, subject="", creator="", producer="")
    return d, frame


def _table(rows, widths, head=True):
    data = [[Paragraph(str(c), CELLB if (head and i == 0) else CELL) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, repeatRows=1 if head else 0)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa0a6")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eaed") if head else colors.white),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 4),
                           ("RIGHTPADDING", (0, 0), (-1, -1), 4)]))
    return t


# ------------------------------------------------------------------------------------------------- Schedule B
def write_schedule_b(path, targets, group_npi):
    title = "Accountable Care Partner Agreement, Schedule B: Annual Reconciliation"
    doc, frame = _doc(path, title, W.FUND, "%s  |  Accountable Care Partner Agreement  |  Schedule B (2026 edition)" % W.FUND)
    s = [Paragraph("ACCOUNTABLE CARE PARTNER AGREEMENT", H1),
         Paragraph("Schedule B: Annual Reconciliation (2026 edition)<br/>Common to the Fund's agreements with each partner "
                   "group listed in Exhibit 1. Adopted by the Board of Trustees on 18 November 2025.", SUB)]
    N = lambda n, t: s.append(Paragraph("<b>%s</b>&nbsp;&nbsp;&nbsp;%s" % (n, t), NUM))
    M = lambda n, t: s.append(Paragraph("(%s)&nbsp;&nbsp;%s" % (n, t), SUBNUM))
    s.append(Paragraph("B.1 Definitions", H2))
    N("B.1.1", "<i>Performance Year</i> means a calendar year. This edition applies to the Performance Year 2026.")
    N("B.1.2", "<i>Member</i> means a participant or dependent eligible for medical benefits under the Fund's Plan of "
               "Benefits in the month concerned.")
    N("B.1.3", "<i>Attributed Member Month</i> means a month in which a Member is eligible for the Fund's medical benefits "
               "and appears on the Fund's attribution roster for the Partner for that month. The roster is run on the "
               "first day of each month and attributes each Member to the partner group of the Member's primary care "
               "provider of record on that day.")
    N("B.1.4", "<i>Total Cost of Care</i> means the Fund's paid amount for covered medical services (professional services "
               "and outpatient and inpatient facility services) with a first date of service in the Performance Year, for "
               "Members attributed to the Partner in the month of that first date of service, whenever the claims are paid. "
               "Outpatient prescription drugs dispensed by retail, mail-order and specialty pharmacies are excluded.")
    N("B.1.5", "Where a Member's amount under B.1.4 for the Performance Year exceeds $200,000, the amount above $200,000 is "
               "excluded. For a Member attributed to more than one partner group during the year, the excluded amount is "
               "shared in proportion to the Member's amounts with each.")
    N("B.1.6", "<i>TCOC PMPM</i> means the Partner's Total Cost of Care for the Performance Year, after B.1.5, divided by "
               "its Attributed Member Months for the year.")
    s.append(Paragraph("B.2 Shared savings and shared losses", H2))
    N("B.2.1", "If the Partner's TCOC PMPM exceeds its Target PMPM in Exhibit 1 by more than 2.0 per cent, the Partner pays "
               "the Fund a shared-loss payment of 40 per cent of the amount by which its TCOC PMPM exceeds its Target PMPM, "
               "multiplied by its Attributed Member Months, but not more than 4 per cent of its Target PMPM multiplied by "
               "its Attributed Member Months.")
    N("B.2.2", "If the Partner's TCOC PMPM is below its Target PMPM by more than 2.0 per cent and its quality score for the "
               "Performance Year is 70 points or more, the Fund pays the Partner a shared-savings payment of 40 per cent "
               "of the amount by which its TCOC PMPM is below its Target PMPM, multiplied by its Attributed Member Months, "
               "on the same limit.")
    N("B.2.3", "The comparisons in B.2.1 and B.2.2 are made on unrounded figures.")
    s.append(Paragraph("B.3 Reconciliation", H2))
    N("B.3.1", "Within 60 days after the end of the Performance Year the Fund issues each Partner a reconciliation from "
               "the claims its claims administrator has paid through the last business day of January following the "
               "year. The Fund estimates the claims incurred in the Performance Year and not yet paid from its claims data.")
    N("B.3.2", "A shared-loss or shared-savings payment under B.2 is due within 45 days after the reconciliation is issued.")
    N("B.3.3", "When twelve months of claims after the Performance Year have been paid, the Fund issues a final "
               "reconciliation on the claims then paid, and any difference from the payment under B.3.2 is settled.")
    s.append(Paragraph("B.4 The reconciliation report", H2))
    N("B.4.1", "For information, the reconciliation also reports for each Partner:")
    M("a", "the Fund's pharmacy spending for its Members: the Fund's net payments to its pharmacy benefit manager for "
           "prescriptions filled in the Performance Year in the Members' Attributed Member Months with the Partner, "
           "reversals included whenever processed; and")
    M("b", "the share of the Fund's paid amount for the professional services in its Total Cost of Care, before B.1.5 "
           "and as paid through the reconciliation date, that was for services rendered by clinicians who were not with "
           "the Partner group when the service was given.")
    s.append(Spacer(1, 6))
    s.append(Paragraph("Exhibit 1: Partner groups and Target PMPM, Performance Year 2026", H2))
    rows = [["Partner group", "Partner code", "Group NPI", "Target PMPM (USD)"]]
    for code in W.CODES:
        rows.append([W.NAME[code], code, group_npi[code], "{:,.2f}".format(targets[code])])
    s.append(_table(rows, [2.6 * inch, 0.9 * inch, 1.2 * inch, 1.4 * inch]))
    s.append(Spacer(1, 6))
    s.append(Paragraph("Target PMPM: as set for 2026 in Schedule A of each Partner's agreement. Group NPI: the "
                       "Partner group's organisational NPI.", SMALL))
    s.append(Spacer(1, 14))
    s.append(Paragraph("For the Board of Trustees: %s, %s." % W.PEOPLE["chair"], BODY))
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ------------------------------------------------------------------------------------------- Keswold dictionary
def write_dictionary(path, n_rows):
    L = []
    a = L.append
    a("KESWOLD BENEFIT SERVICES")
    a("Client extract: medical claims")
    a("Client: %s (group 4417)" % W.FUND)
    a("Extract: medical_claims_extract_2025-01_2027-01.csv")
    a("Run: 1 February 2027 05:12 (job CLX-MED-0221), %d rows" % n_rows)
    a("Selection: every version of every medical claim with a first date of service from 1 January 2025 to")
    a("31 January 2027, as adjudicated through the payment run of 29 January 2027.")
    a("")
    a("One row per claim version. A claim keeps its claim_number through every version.")
    a("")
    a("FIELDS")
    flds = [
        ("claim_number", "assigned on receipt: two-digit year, day of the year received, daily sequence"),
        ("claim_version", "1 for the claim as first adjudicated, 2 for a later replacement or void"),
        ("claim_frequency_code", "the version's type (see CODES)"),
        ("member_id", "the Fund's member id"),
        ("claim_form", "CMS1500 (professional) or UB04 (facility)"),
        ("bill_type", "UB-04 type of bill for facility claims"),
        ("place_of_service", "CMS place of service code for professional claims"),
        ("service_from_date", "first date of service on the claim"),
        ("service_to_date", "last date of service on the claim (the discharge date for an inpatient stay)"),
        ("admission_date", "inpatient claims only"),
        ("received_date", "the date Keswold received the claim (electronic or paper)"),
        ("paid_date", "the date of the payment run on which this version was paid or recovered"),
        ("billed_amount", "provider's billed charges, USD"),
        ("allowed_amount", "allowed amount after network pricing, USD"),
        ("paid_amount", "the Fund's paid amount for the claim as of this version, USD"),
        ("payment_amount", "the amount paid (positive) or recovered (negative) with this version, USD"),
        ("claim_status", "P paid, D denied, V voided"),
        ("rendering_npi", "the clinician who rendered the service (professional claims)"),
        ("billing_npi", "the billing provider: a group or facility NPI, or the clinician's own"),
    ]
    for n, d in flds:
        a("  %-21s %s" % (n, d))
    a("")
    a("CODES")
    a("  claim_frequency_code")
    a("    Claims adjudicated on the Tallis platform:   1 original, 7 replacement, 8 void")
    a("    Claims adjudicated on the legacy system:     O original, A adjustment (replacement), V void")
    a("  A replacement carries the claim's full corrected amounts; a void carries zero amounts. The claim's")
    a("  value is the paid_amount of its latest version.")
    a("")
    a("RELEASE NOTES")
    a("  2026.10  5 October 2026. The Fund's claims are adjudicated on the Tallis platform. Versions processed")
    a("           from this date carry the Tallis frequency codes; claim numbers are unchanged. Claims received")
    a("           before 5 October and not yet paid were adjudicated on Tallis from that date.")
    a("  2025.04  7 April 2025. place_of_service added for professional claims.")
    a("")
    a("Contact: %s, %s" % W.PEOPLE["tpa"])
    open(path, "w", encoding="utf-8", newline="\r\n").write("\n".join(L) + "\n")


# ----------------------------------------------------------------------------------------- the Fund's data notes
def write_notes(path):
    import docx
    from docx.shared import Pt
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    d.add_heading("Laurelford Fund data notes", level=1)
    d.add_paragraph("Fund Office, February 2027. For analysts using the Fund's eligibility and attribution files. "
                    "Questions to %s, %s." % W.PEOPLE["administrator"])
    d.add_heading("Member ids", level=2)
    d.add_paragraph("Each participant has a subscriber id (LB and seven digits). Every member of the family, the "
                    "participant included, has a member id made of the subscriber id and a two-digit suffix. Suffixes "
                    "are assigned in the order the Fund enrolls the family's members, starting with 01 for the "
                    "participant, and a suffix is never reused.")
    d.add_heading("Eligibility spans (eligibility_spans_2025-2027.csv)", level=2)
    d.add_paragraph("Medical eligibility is earned through the hour bank: the hours an employer reports for a work month, "
                    "with the participant's banked hours, decide coverage two months later. The file holds one row per "
                    "version of each coverage span. A later version of a span (span_version 2) replaces the earlier one, "
                    "for example when an employer's contributions are reversed after an audit (change_reason RT). When "
                    "late contributions post for a month the participant was not covered, the month is reinstated as a "
                    "span of its own (change_reason RI). coverage_end is blank for a span still open on 31 January 2027. The "
                    "file is the export of 1 February 2027 and holds the versions posted by then.")
    d.add_paragraph("relationship_code: P participant, S spouse, D dependent child. change_reason: NE new enrollee, "
                    "HB hour bank, RI reinstated after late contributions, RT retroactive termination.")
    d.add_heading("Attribution roster (attribution_roster_2026.csv)", level=2)
    d.add_paragraph("The roster is run on the first day of each month from the eligibility on file that day. It lists "
                    "each member whose primary care provider of record belongs to an accountable care partner group, "
                    "with that provider's NPI and the group's partner code. Members without a partner primary care "
                    "provider do not appear. The roster is not rerun when eligibility changes later.")
    d.add_heading("Other files", level=2)
    d.add_paragraph("Pharmacy claims come from Corvane Pharmacy Benefits in Corvane's own layout. The partners' clinician "
                    "roster is the network file the partners update by notice.")
    d.save(path)


# --------------------------------------------------------------------------------------------- Corvane layout
def write_rx_layout(path):
    title = "Client claims file layout, version 4"
    doc, frame = _doc(path, title, W.PBM, "%s  |  Client claims file layout v4  |  Confidential" % W.PBM)
    s = [Paragraph("CORVANE PHARMACY BENEFITS", H1),
         Paragraph("Client claims file layout, version 4 (effective 1 January 2025)", SUB),
         Paragraph("The client claims file carries every pharmacy claim transaction adjudicated for the client's plan in "
                   "the period, one row per transaction, comma-separated with a header row. Amounts are in US dollars.", BODY)]
    rows = [["Field", "Description"],
            ["claim_id", "Corvane's transaction id"],
            ["cardholder_id", "the id on the member's card: the plan's subscriber id for the family"],
            ["person_code", "identifies the patient within the cardholder's family: 01 the cardholder, 02 the spouse, "
                            "03 and up the dependents in order of date of birth, eldest first"],
            ["patient_birth_date", "the patient's date of birth"],
            ["patient_gender", "M or F"],
            ["relationship_code", "1 cardholder, 2 spouse, 3 dependent"],
            ["rx_number", "the pharmacy's prescription number"],
            ["fill_date", "the date the prescription was filled"],
            ["ndc", "National Drug Code (5-4-2)"],
            ["quantity, days_supply", "as dispensed"],
            ["ingredient_cost, dispensing_fee", "the plan's price for the drug and the pharmacy's fee"],
            ["patient_pay", "the member's copay or coinsurance"],
            ["plan_paid", "the amount the plan pays for the transaction"],
            ["transaction_code", "B1 claim, B2 reversal, B3 rebill (NCPDP). A reversal carries the amounts of the claim "
                                 "it reverses; a rebill replaces a reversed claim"],
            ["adjudication_date", "the date Corvane adjudicated the transaction"]]
    s.append(_table(rows, [1.9 * inch, 4.6 * inch]))
    s.append(Spacer(1, 8))
    s.append(Paragraph("Corvane invoices the client twice a month for plan_paid on the transactions adjudicated in the "
                       "billing period, net of reversals.", BODY))
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# --------------------------------------------------------------------------------------------- stop-loss summary
def write_stop_loss(path):
    title = "Specific stop-loss policy summary, 2026"
    doc, frame = _doc(path, title, W.FUND, "%s  |  Stop-loss policy summary 2026" % W.FUND)
    s = [Paragraph("SPECIFIC STOP-LOSS INSURANCE: POLICY SUMMARY", H1),
         Paragraph("Policy year 1 January to 31 December 2026. Carrier: %s. Policyholder: %s." % (W.STOP_LOSS_CARRIER, W.FUND), SUB)]
    rows = [["Term", "Provision"],
            ["Specific deductible", "$175,000 per covered person per policy year"],
            ["Contract basis", "12/15: claims incurred in the policy year and paid from 1 January 2026 to 31 March 2027"],
            ["Covered expenses", "medical claims and outpatient prescription drugs paid by the Plan"],
            ["Reimbursement", "100 per cent of covered expenses above the specific deductible, on submission of proof "
                              "of payment"],
            ["Aggregate coverage", "none"]]
    s.append(_table(rows, [1.8 * inch, 4.7 * inch]))
    s.append(Spacer(1, 8))
    s.append(Paragraph("Reimbursements are credited to the Plan's claims account when received. This summary does not "
                       "change the policy's terms.", BODY))
    s.append(Paragraph("Prepared by %s, %s, December 2025." % W.PEOPLE["controller"], BODY))
    doc.build(s, onFirstPage=frame, onLaterPages=frame)
