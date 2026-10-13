"""Write the Laurelford pack's data files from a simulated world: the administrator's claims extract, the Fund's
eligibility spans and attribution roster, Corvane's pharmacy file, the partners' clinician roster and its 2026
changes, and the quality scorecard."""
import csv
import datetime as dt

import numpy as np

import sim as S
import world as W

FILES = {
    "claims": "medical_claims_extract_2025-01_2027-01.csv",
    "dictionary": "keswold_extract_dictionary_2027-02-01.txt",
    "elig": "eligibility_spans_2025-2027.csv",
    "roster": "attribution_roster_2026.csv",
    "notes": "laurelford_data_notes_2027-02.docx",
    "agreement": "accountable_care_partner_agreement_schedule_b.pdf",
    "rx": "corvane_pharmacy_claims_2026.csv",
    "rx_layout": "corvane_claims_file_layout_v4.pdf",
    "clinicians": "partner_clinician_roster_2027-01-01.xlsx",
    "stoploss": "stop_loss_policy_2026.pdf",
    "quality": "partner_quality_scorecard_2026.xlsx",
}
DISTRACTORS = ("stoploss", "quality")

SPECIALTIES = ["Cardiology", "Orthopaedic Surgery", "Gastroenterology", "Dermatology", "Endocrinology", "Pulmonology",
               "Neurology", "Obstetrics and Gynecology", "General Surgery", "Urology", "Psychiatry", "Rheumatology"]
PCP_SPECIALTY = {"adult": ["Internal Medicine", "Family Medicine"], "family": ["Family Medicine"],
                 "pediatric": ["Pediatrics"]}
POS = {"office": "11", "outpatient": "22", "er": "23", "telehealth": "10"}


def _w(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh, lineterminator="\r\n")
        wr.writerow(header)
        wr.writerows(rows)
    return len(rows)


def m2(x):
    return "%.2f" % x


def claim_numbers(world):
    """the administrator numbers a claim on receipt: two-digit year, Julian day, then a daily sequence"""
    seq = {}
    out = {}
    for c in sorted(world["claims"], key=lambda c: (c["recv"], c["seq"])):
        n = seq.get(c["recv"], 0) + 1
        seq[c["recv"]] = n
        out[c["seq"]] = "%02d%03d%05d" % (c["recv"].year % 100, c["recv"].timetuple().tm_yday, n)
    return out


def write_claims(world, path):
    nums = claim_numbers(world)
    rng = np.random.default_rng(W.SEED + 11)
    rows = []
    for v in sorted(world["versions"], key=lambda v: (v["vpaid"], v["seq"], v["ver"])):
        if v["vpaid"] > W.CUT or v["sfrom"] < W.SERVICE_FROM or v["sfrom"] > W.SERVICE_TO:
            continue
        legacy = v["vsys"] == "L"
        code = {"orig": "O" if legacy else "1", "repl": "A" if legacy else "7", "void": "V" if legacy else "8"}[v["kind"]]
        prev = 0.0 if v["ver"] == 1 else v["amt"]
        status = "D" if v["deny"] else ("V" if v["kind"] == "void" else "P")
        t = v["t"]
        form = "CMS1500" if t == "P" else "UB04"
        bill_type = {"P": "", "O": "0131", "I": "0111"}[t]
        pos = ""
        if t == "P":
            u = (v["seq"] * 2654435761) % 1000 / 1000.0
            pos = POS["office"] if u < 0.78 else (POS["outpatient"] if u < 0.88 else (POS["er"] if u < 0.95 else POS["telehealth"]))
        rows.append([nums[v["seq"]], v["ver"], code, v["mid"], form, bill_type, pos, v["sfrom"].isoformat(),
                     v["sto"].isoformat(), v["admit"].isoformat() if v["admit"] else "", v["recv"].isoformat(),
                     v["vpaid"].isoformat(), m2(v["billed"]), m2(v["vallowed"]), m2(v["vamt"]), m2(v["vamt"] - prev),
                     status, v["rendering"] or "", v["billing"] or (v["rendering"] or "")])
    hdr = ["claim_number", "claim_version", "claim_frequency_code", "member_id", "claim_form", "bill_type",
           "place_of_service", "service_from_date", "service_to_date", "admission_date", "received_date", "paid_date",
           "billed_amount", "allowed_amount", "paid_amount", "payment_amount", "claim_status", "rendering_npi",
           "billing_npi"]
    return _w(path, hdr, rows)


def write_elig(world, path):
    """the export of 1 February 2027: the span versions posted by then (a later reinstatement or termination is not
    yet on file)"""
    rows = sorted((r for r in world["elig"] if r[9] <= W.ELIG_EXPORT.isoformat()), key=lambda r: (r[1], r[0], r[5], r[6]))
    hdr = ["member_id", "subscriber_id", "relationship_code", "birth_date", "gender", "span_id", "span_version",
           "coverage_start", "coverage_end", "posted_date", "change_reason"]
    return _w(path, hdr, rows)


def write_roster(world, path):
    rows = sorted(world["roster"], key=lambda r: (r[0], r[3], r[1]))
    return _w(path, ["roster_month", "member_id", "pcp_npi", "partner_code"], rows)


def write_rx(world, path):
    rows = []
    cid = 0
    for r in sorted(world["rx"], key=lambda r: (r["adj"], r["rx"], r["tc"])):
        cid += 1
        p = r["p"]
        rel = {"P": "1", "S": "2", "D": "3"}[p["rel"]]
        rows.append(["CRV%09d" % (40000000 + cid * 7), p["sid"], p["pc"], p["dob"].isoformat(), p["gender"], rel,
                     r["rx"], r["fill"].isoformat(), r["ndc"], "%g" % r["qty"], r["ds"], m2(r["ing"]), m2(r["fee"]),
                     m2(r["copay"]), m2(r["plan"]), r["tc"], r["adj"].isoformat()])
    hdr = ["claim_id", "cardholder_id", "person_code", "patient_birth_date", "patient_gender", "relationship_code",
           "rx_number", "fill_date", "ndc", "quantity", "days_supply", "ingredient_cost", "dispensing_fee",
           "patient_pay", "plan_paid", "transaction_code", "adjudication_date"]
    return _w(path, hdr, rows)


def clinician_rows(world):
    C = world["C"]
    rng = np.random.default_rng(W.SEED + 23)
    rows = []
    moved = {m["npi"]: m for m in C["moves"]}
    for code in W.CODES:
        for role, pool in (("PCP", C["pcp"][code]), ("Specialist", C["spec"][code])):
            for n in pool:
                m = moved.get(n)
                if m and m["to"] is None:          # left during 2026: not on the January roster
                    continue
                if m and m["to"] == code:
                    eff = m["date"]
                else:
                    eff = S.mstart(int(rng.integers(S.mkey(dt.date(2016, 1, 1)), S.mkey(dt.date(2025, 1, 1)))))
                spec = (PCP_SPECIALTY[W.KIND[code]][int(rng.integers(0, len(PCP_SPECIALTY[W.KIND[code]])))]
                        if role == "PCP" else SPECIALTIES[int(rng.integers(0, len(SPECIALTIES)))])
                rows.append([n, code, W.NAME[code], role, spec, eff])
    rows.sort(key=lambda r: (r[1], r[3], r[0]))
    return rows


def change_rows(world):
    C = world["C"]
    rows = []
    for m in C["moves"]:
        if m["frm"]:
            rows.append([m["npi"], "TERM", m["frm"], (m["date"] - dt.timedelta(days=1)).isoformat(),
                         (m["date"] - dt.timedelta(days=12)).isoformat()])
        if m["to"]:
            rows.append([m["npi"], "JOIN", m["to"], m["date"].isoformat(), (m["date"] - dt.timedelta(days=9)).isoformat()])
    rows.sort(key=lambda r: (r[3], r[0], r[1]))
    return rows


def write_changes(world, path):
    return _w(path, ["npi", "change", "partner_code", "effective_date", "notified_date"], change_rows(world))


def write_clinicians(world, path):
    import openpyxl
    from openpyxl.styles import Alignment, Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clinicians"
    ws["A1"] = "Accountable care partners: clinician roster"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Clinicians in each partner group on 1 January 2027, as the partners notified the Fund. A clinician who "
                "leaves a group is removed when the partner gives notice. Source: Fund network file.")
    hdr = ["NPI", "Partner code", "Partner", "Role", "Specialty", "With partner since"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(bold=True)
    rows = clinician_rows(world)
    for i, r in enumerate(rows, 5):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            if j == 6:
                c.number_format = "mm/dd/yyyy"
    for col, w in zip("ABCDEF", [13, 12, 32, 11, 26, 18]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    wb.save(path)
    return len(rows)


def write_quality(world, path):
    """the partners' 2026 quality scorecard (a distractor: the agreement's quality gate applies to shared savings)"""
    import openpyxl
    from openpyxl.styles import Font
    rng = np.random.default_rng(W.SEED + 31)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2026 measures"
    ws["A1"] = "Accountable care partners: 2026 quality measures"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Preliminary rates, measurement year 2026, from claims paid through 31 December 2026 and practice reports."
    hdr = ["Partner", "Breast cancer screening", "Colorectal cancer screening", "Controlling high blood pressure",
           "Diabetes HbA1c control (<8.0%)", "Childhood immunisation status (Combo 10)", "Quality score (points of 100)"]
    for j, h in enumerate(hdr, 1):
        ws.cell(row=4, column=j, value=h).font = Font(bold=True)
    for i, code in enumerate(W.CODES, 5):
        ws.cell(row=i, column=1, value=W.NAME[code])
        vals = []
        for j in range(5):
            if W.KIND[code] == "pediatric" and j < 4:
                ws.cell(row=i, column=2 + j, value="n/a")
                continue
            if W.KIND[code] == "adult" and j == 4:
                ws.cell(row=i, column=2 + j, value="n/a")
                continue
            v = float(np.clip(rng.normal(0.66, 0.08), 0.42, 0.88))
            vals.append(v)
            c = ws.cell(row=i, column=2 + j, value=round(v, 3))
            c.number_format = "0.0%"
        ws.cell(row=i, column=7, value=round(float(np.mean(vals)) * 100 + rng.normal(4, 3), 1))
    for col, w in zip("ABCDEFG", [32, 14, 14, 16, 16, 18, 14]):
        ws.column_dimensions[col].width = w
    wb.save(path)
    return len(W.CODES)
