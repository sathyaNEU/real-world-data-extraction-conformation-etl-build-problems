"""The golden deliverables, computed from the shipped files only: partner_settlement_2026.xlsx (the Fund Office's 2026
settlement workbook: each partner's attributed member months, completed total cost of care per member month against its
target, whether it owes, its members' 2026 pharmacy spending and the share of its members' professional spending that
went to clinicians outside the group) and partners_against_target_2026.svg (each partner's cost as a share of its target,
the corridor marked); and figures.json, every figure the write-up states."""
import collections
import datetime as dt
import io
import json
import os
import re
import sys

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import calc as C        # noqa: E402
import pack as P        # noqa: E402
import world as W       # noqa: E402

XLSX = "partner_settlement_2026.xlsx"
PNG = "partners_against_target_2026.svg"     # the chart deliverable
GOLD = "G1 received, all development"
FAMILY = ("G1 received, all development", "G2 received, latest 12 months", "G3 paid, after the go-live month",
          "G4 received, by claim type", "G5 paid (latest versions), after the go-live month")
RECEIVED = ("G1 received, all development", "G2 received, latest 12 months", "G4 received, by claim type")
ROUTES = ("R0 paid, all development", "paid, latest 12 months", "paid, latest 6 months", "paid, latest 4 months",
          "R1 paid as extracted")


def exhibit(target):
    """Exhibit 1 of Schedule B: each partner's group NPI and Target PMPM"""
    from pypdf import PdfReader
    txt = "\n".join(p.extract_text() for p in PdfReader(os.path.join(target, P.FILES["agreement"])).pages)
    tg, gn = {}, {}
    for code in W.CODES:
        m = re.search(r"\b%s\s+(\d{10})\s+([\d,]+\.\d\d)" % code, txt)
        gn[code], tg[code] = m.group(1), float(m.group(2).replace(",", ""))
    return tg, gn


def lags(cl):
    """the forcing evidence: days from receipt to payment and from service to receipt, before and after go-live"""
    v1 = cl[cl["claim_version"] == 1].copy()
    rcv = v1["received_date"].map(dt.date.fromisoformat)
    pdt = v1["paid_date"].map(dt.date.fromisoformat)
    sto = v1["service_to_date"].map(dt.date.fromisoformat)
    r2p = np.array([(b - a).days for a, b in zip(rcv, pdt)])
    s2r = np.array([(b - a).days for a, b in zip(sto, rcv)])
    before = np.array([a < W.GOLIVE for a in rcv])
    after = np.array([a >= W.GOLIVE for a in rcv])
    pre_year = np.array([dt.date(2025, 1, 1) <= a < dt.date(2026, 1, 1) for a in rcv])
    post = after
    return dict(r2p_median_before=float(np.median(r2p[before])), r2p_median_after=float(np.median(r2p[after])),
                r2p_mean_before=float(np.mean(r2p[before])), r2p_mean_after=float(np.mean(r2p[after])),
                s2r_median_2025=float(np.median(s2r[pre_year])), s2r_median_after=float(np.median(s2r[post])),
                s2r_mean_2025=float(np.mean(s2r[pre_year])), s2r_mean_after=float(np.mean(s2r[post])))


def compute(target):
    tg, gn = exhibit(target)
    X = C.compute(target, tg)
    X["targets"], X["group_npi"] = tg, gn
    X["rx"] = C.ask_pharmacy(X["L"], X["att"], "dob", X["elig"])
    X["share"] = C.ask_leakage(X["L"], X["d26"], gn, "billing")
    X["lags"] = lags(X["L"]["claims"])
    X["call"] = X["flags"][GOLD]
    return X


def k1(v):
    return int(round(v / 1000.0)) * 1000


def figures(X):
    F = collections.OrderedDict()
    tg = X["targets"]
    g = X["pmpm"][GOLD]
    F["call"] = [W.NAME[c] for c in W.CODES if c in X["call"]]
    F["call_codes"] = [c for c in W.CODES if c in X["call"]]
    F["n_owe"] = len(X["call"])
    F["routes"] = collections.OrderedDict((nm, [c for c in W.CODES if c in X["flags"][nm]]) for nm in X["flags"])
    F["default_n"] = len(X["flags"]["R0 paid, all development"])
    F["ratio_pct"] = collections.OrderedDict(
        (nm, collections.OrderedDict((c, round(100 * X["pmpm"][nm][c] / tg[c], 2)) for c in W.CODES)) for nm in X["pmpm"])
    fam = {c: [100 * X["pmpm"][nm][c] / tg[c] for nm in FAMILY] for c in W.CODES}
    F["family_ratio_range_pct"] = collections.OrderedDict((c, ["%.1f" % min(v), "%.1f" % max(v)]) for c, v in fam.items())
    rec = {c: [100 * X["pmpm"][nm][c] / tg[c] for nm in RECEIVED] for c in W.CODES}
    F["received_ratio_range_pct"] = collections.OrderedDict((c, ["%.1f" % min(v), "%.1f" % max(v)]) for c, v in rec.items())
    F["golden_tcoc_pmpm"] = collections.OrderedDict((c, round(g[c], 2)) for c in W.CODES)
    F["targets"] = collections.OrderedDict((c, tg[c]) for c in W.CODES)
    F["member_months"] = collections.OrderedDict((c, int(X["mm"][c])) for c in W.CODES)
    F["truncated_members"] = len(X["cats"])
    F["agreement"] = collections.OrderedDict([("truncation", W.TRUNCATION), ("corridor_pct", 100 * W.CORRIDOR), ("line_pct", 100 + 100 * W.CORRIDOR)])
    F["pharmacy"] = collections.OrderedDict((c, round(X["rx"][c], 2)) for c in W.CODES)
    F["pharmacy_stated"] = collections.OrderedDict((c, k1(X["rx"][c])) for c in W.CODES)
    F["outside_share"] = collections.OrderedDict((c, round(X["share"][c], 4)) for c in W.CODES)
    F["outside_share_stated"] = collections.OrderedDict((c, "%.1f" % X["share"][c]) for c in W.CODES)
    lg = X["lags"]
    F["lags"] = collections.OrderedDict((k, round(v, 1)) for k, v in lg.items())
    i0 = W.YEAR * 12 - C.S0
    book = {nm: sum(X["pmpm"][nm][c] * X["mm"][c] for c in W.CODES) for nm in X["pmpm"]}
    F["book_vs_golden_pct"] = collections.OrderedDict((nm, round(100 * (book[nm] / book[GOLD] - 1), 2)) for nm in X["pmpm"])
    F["book_vs_golden_stated"] = collections.OrderedDict((nm, "%.0f" % (100 * (book[nm] / book[GOLD] - 1))) for nm in X["pmpm"])
    return F


# ------------------------------------------------------------------------------------------------------- chart
def chart_png(X, F):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.family": "DejaVu Sans", "svg.fonttype": "none", "svg.hashsalt": "task451"})
    tg = X["targets"]
    g = X["pmpm"][GOLD]
    order = sorted(W.CODES, key=lambda c: g[c] / tg[c])
    vals = [100 * g[c] / tg[c] for c in order]
    owe = set(X["call"])
    col = ["#a33b20" if c in owe else "#8c9aa6" for c in order]
    fig, ax = plt.subplots(figsize=(8.8, 5.4), dpi=150)
    ax.barh([W.NAME[c] for c in order], vals, color=col, height=0.62)
    ax.axvline(100, color="#30363d", lw=1.0)
    ax.axvline(102, color="#a33b20", lw=1.0, ls="--")
    ax.set_ylim(-0.6, len(order) + 0.25)
    ax.text(102.3, len(order) - 0.2, "Shared loss above 102% of target", fontsize=8, color="#a33b20", va="center")
    for i, v in enumerate(vals):
        ax.text(v + 0.4, i, "%.1f%%" % v, va="center", fontsize=8.2, color="#202326")
    ax.set_xlim(min(vals) - 6, max(vals) + 6)
    ax.set_xlabel("2026 total cost of care per attributed member month, as a share of the partner's target (%)", fontsize=8.8)
    ax.tick_params(axis="both", labelsize=8.4)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    names = [" ".join(W.NAME[c].split()[:-2]) for c in reversed(order) if c in owe]
    words = {2: "Two", 3: "Three", 4: "Four", 5: "Five", 6: "Six"}
    ax.set_title("%s partners owe a 2026 shared-loss payment\n%s and %s are more than 2 per cent over target" % (
        words.get(len(names), str(len(names))), ", ".join(names[:-1]), names[-1]), fontsize=9.6, loc="left", color="#202326")
    fig.text(0.01, 0.01, "Claims paid through 29 January 2027; claims not yet paid estimated from the claims' received dates. "
                         "Laurelford Fund Office.", fontsize=7, color="#5a5f66")
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    buf = io.BytesIO()
    fig.savefig(buf, format="svg", metadata={"Date": None, "Creator": None, "Format": None, "Type": None})
    plt.close(fig)
    import re
    s = buf.getvalue().decode("utf-8")
    s = re.sub(r"<!-- Created with [^>]*-->\n?", "", s)
    s = re.sub(r" *<metadata>.*?</metadata>\n?", "", s, flags=re.S)
    s = s.replace('id="matplotlib.axis_', 'id="axis_')
    assert "matplotlib" not in s.lower(), "plotting library named in the SVG"
    return s.encode("utf-8")


# ------------------------------------------------------------------------------------------------------ workbook
def write_xlsx(X, F, path):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    tg = X["targets"]
    g = X["pmpm"][GOLD]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Settlement 2026"
    bold = Font(bold=True)
    ws["A1"] = "Accountable care partners: 2026 reconciliation (Schedule B)"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Claims paid through 29 January 2027. Claims incurred in 2026 and not yet paid are estimated from the "
                "claims' received dates (see Notes). For the Board of Trustees, 17 February 2027.")
    hdr = ["Partner group", "Code", "Attributed member months", "Total cost of care PMPM ($)", "Target PMPM ($)",
           "Share of target", "Owes a shared-loss payment", "Pharmacy spending ($)", "Professional spending outside the group"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = bold
        c.fill = PatternFill("solid", fgColor="E3E8EC")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, code in enumerate(W.CODES, 5):
        vals = [W.NAME[code], code, X["mm"][code], round(g[code], 2), tg[code], round(g[code] / tg[code], 4),
                "Yes" if code in X["call"] else "No", F["pharmacy_stated"][code], round(X["share"][code] / 100, 3)]
        fmts = [None, None, "#,##0", "#,##0.00", "#,##0.00", "0.0%", None, "#,##0", "0.0%"]
        for j, (v, fm) in enumerate(zip(vals, fmts), 1):
            c = ws.cell(row=i, column=j, value=v)
            if fm:
                c.number_format = fm
            if j == 7:
                c.alignment = Alignment(horizontal="center")
        if code in X["call"]:
            for j in range(1, 10):
                ws.cell(row=i, column=j).font = bold
    r = 5 + len(W.CODES)
    ws.cell(row=r, column=1, value="All partners").font = bold
    ws.cell(row=r, column=3, value=int(sum(X["mm"][c] for c in W.CODES))).number_format = "#,##0"
    ws.cell(row=r, column=8, value=k1(sum(X["rx"][c] for c in W.CODES))).number_format = "#,##0"
    for j in range(1, 10):
        ws.cell(row=r, column=j).border = Border(top=Side(style="thin"))
    ws.cell(row=r + 2, column=1, value=("A partner owes if its total cost of care per member month is more than 2.0 per "
                                         "cent over its target (Schedule B, B.2.1), compared before rounding. Pharmacy "
                                         "spending to the nearest thousand dollars, the total rounded from the unrounded amounts; shares to one decimal "
                                         "place."))
    ws.cell(row=r + 2, column=1).font = Font(italic=True, color="5A5F66")
    for col, w in zip("ABCDEFGHI", [31, 7, 13, 15, 12, 10, 13, 14, 16]):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[4].height = 46
    ws.freeze_panes = "C5"
    ws.print_title_rows = "4:4"
    ws.print_area = "A1:I%d" % (r + 2)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    nt = wb.create_sheet("Notes")
    nt["A1"] = "Notes"
    nt["A1"].font = bold
    lg = X["lags"]
    notes = [
        "Total cost of care (Schedule B, B.1.4 to B.1.6): the Fund's paid amount for medical claims with a first date of "
        "service in 2026, each claim at its latest version, for members attributed to the partner in the month of that "
        "first date of service; each member's year above $200,000 excluded; divided by attributed member months (the "
        "monthly roster, for months in which the member's eligibility stands on the latest version of the eligibility "
        "spans).",
        "Claims not yet paid on 29 January 2027 (B.3.1). Keswold moved the Fund's claims to the Tallis platform on 5 "
        "October 2026, and the time from receipt to payment fell from a median of %d days to %d. The time from service "
        "to receipt, which the providers set, did not change (a median of %d days in 2025 and %d since). The 2026 claims "
        "are therefore completed on their received dates: each month's claims received by the cut, grossed up by how "
        "much of a month's claims had been received at the same point in earlier months, all of which Tallis has now "
        "paid within days." % (lg["r2p_median_before"], lg["r2p_median_after"], lg["s2r_median_2025"], lg["s2r_median_after"]),
        "Completing on paid dates instead would carry the old payment speed into October to December and overstate "
        "the year's cost by about %s per cent." % F["book_vs_golden_stated"]["R0 paid, all development"],
        "Pharmacy spending (B.4.1(a)): Corvane's plan paid amounts for prescriptions filled in 2026, net of reversals "
        "whenever processed, "
        "in the member's attributed member months with the partner (on the roster for the month of the fill and "
        "eligible in it on the latest eligibility spans). Corvane's person codes number dependants by date of birth and "
        "the Fund's member ids by order of enrollment, so each patient is matched to the member with the same "
        "cardholder, date of birth and sex.",
        "Professional spending outside the group (B.4.1(b)): of the Fund's paid amount for the professional claims in "
        "the total cost of care, before the $200,000 exclusion and as paid through 29 January 2027, the share billed by "
        "anyone other than the partner group (its group NPI in Exhibit 1): the group a clinician billed through when "
        "the service was given.",
        "Sources: Keswold claims extract of 1 February 2027; Fund eligibility spans and 2026 attribution roster; "
        "Corvane client claims file 2026; Schedule B and Exhibit 1.",
    ]
    for i, t in enumerate(notes, 3):
        nt.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True, vertical="top")
    nt.column_dimensions["A"].width = 120
    wb.save(path)


def main(target, golden, figures_path):
    os.makedirs(golden, exist_ok=True)
    X = compute(target)
    F = figures(X)
    write_xlsx(X, F, os.path.join(golden, XLSX))
    open(os.path.join(golden, PNG), "wb").write(chart_png(X, F))
    import postprocess as PP
    by = W.PEOPLE["administrator"][0]
    PP.normalize_ooxml(os.path.join(golden, XLSX), by, by, "2027-02-03T14:05:00Z", "2027-02-04T22:16:43Z", excel=True,
                       company=W.FUND)
    with open(figures_path, "w") as fh:
        json.dump(F, fh, indent=1, ensure_ascii=False)
        fh.write("\n")
    return X
