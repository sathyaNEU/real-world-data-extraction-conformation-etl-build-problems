"""task172 independent verifier: reads only the shipped bundle (target/), never the generator's world or calc.py, and
recomputes the quote by exact convolution rather than simulation, then both ask tables, and compares every owed figure
with generator/figures.json at the prompt's rounding.

For a renewal of the coming quarter the natural wait is the appointment lead (with its category's appointment share)
plus the review; once the natural wait has passed the team's request (the card's expiry plus the request offset), an
accepted request ends it at the request plus the expedite response. So P(W > w) = P(natural > w) * (1 - a * P(gap +
offset + response <= w)), every distribution read from the record in whole days, the natural wait mixed over the
quarter's categories.

    python3 task172/generator/verify.py task172/target"""
from __future__ import annotations

import csv
import datetime as dt
import json
import os
import sys
from collections import defaultdict

import numpy as np
from openpyxl import load_workbook

EXPORT = dt.date(2026, 12, 4)
RULE = dt.date(2025, 10, 30)
LEAD_WINDOW = (dt.date(2026, 10, 1), dt.date(2026, 11, 30))
QUARTER = (dt.date(2027, 1, 1), dt.date(2027, 3, 31))
MONTH = 30.4
FAILS = []


def say(ok, label, detail=""):
    print("%s  %-80s %s" % ("PASS" if ok else "FAIL", label, detail))
    if not ok:
        FAILS.append(label)


def day(v):
    if v in (None, ""):
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    v = str(v)
    if "/" in v:
        m, d, y = v.split("/")
        return dt.date(int(y), int(m), int(d))
    return dt.date.fromisoformat(v[:10])


def rows_csv(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def rows_xlsx(path, sheet):
    ws = load_workbook(path, read_only=True)[sheet]
    it = ws.iter_rows(values_only=True)
    head = next(it)
    return [dict(zip(head, r)) for r in it]


def kaplan_meier(times, events):
    """survival just after each whole day 0..T for (time, event) data"""
    t = np.asarray(times, int)
    e = np.asarray(events, int)
    T = int(t.max()) + 1
    deaths = np.bincount(t[e == 1], minlength=T + 1)
    leaving = np.bincount(t, minlength=T + 1)
    at_risk = len(t) - np.concatenate([[0], np.cumsum(leaving)[:-1]])
    S = np.cumprod(np.where(at_risk > 0, 1 - deaths / np.maximum(at_risk, 1), 1.0))
    return S                        # S[k] = P(Q > k)


def main(target):
    root = os.path.dirname(os.path.abspath(target))
    figs = json.load(open(os.path.join(root, "generator", "figures.json")))
    notices = rows_csv(os.path.join(target, "renewal_case_notices_2024-01_2026-12.csv"))
    actions = rows_csv(os.path.join(target, "case_actions_2024-2026.csv"))
    reg_path = os.path.join(target, "renewal_case_register.xlsx")
    old = rows_xlsx(reg_path, "Cases to Mar 2025")
    new = rows_xlsx(reg_path, "Cases")
    # one key per case: the current reference, reached from either reference
    key = {}
    case = {}
    for r in old:
        key[r["Case ID"]] = r["Case ID"]
        case[r["Case ID"]] = dict(person=r["Person ID"], ptype=r["Person Type"], sponsor=r["Sponsor ID"], cat=r["Category"],
                                  kind=r["Case Type"].lower(), card=day(r["Card Expires"]), received=day(r["Received"]),
                                  status=r["Status"], closed=day(r["Closed"]), decided=day(r["Decided"]),
                                  new=day(r["New Card Expires"]))
    for r in new:
        key[r["case_id"]] = r["case_id"]
        if r["legacy_case_id"]:
            key[r["legacy_case_id"]] = r["case_id"]
        case[r["case_id"]] = dict(person=r["person_id"], ptype=r["person_type"].title(), sponsor=r["sponsor_id"],
                                  cat=r["category"], kind=r["case_type"], card=day(r["card_expires"]),
                                  received=day(r["received_on"]), status=r["status"], closed=day(r["closed_on"]),
                                  decided=day(r["decided_on"]), new=day(r["new_card_expires"]))
    # renewals by receipt number
    by = defaultdict(list)
    for n in notices:
        if n["receipt_no"]:
            by[n["receipt_no"]].append(n)
    req = defaultdict(list)
    stop = {}
    for a in actions:
        k = key.get(a["case_id"])
        if a["action"] == "Expedite request" and k:
            req[k].append((day(a["action_date"]), a["outcome"]))
        if a["action"] == "Stop work" and k:
            stop[k] = day(a["action_date"])
    ren = []
    for rno, ns in by.items():
        rc = [n for n in ns if n["notice_type"] == "Receipt notice"]
        if not rc:
            continue
        k = key.get(ns[0]["client_ref"])
        rcv = day(rc[0]["received_date"])
        ap = sorted([n for n in ns if n["notice_type"] == "Biometrics appointment notice"],
                    key=lambda n: (n["notice_date"], n["logged_on"], n["notice_id"]))
        reuse = any(n["notice_type"] == "Biometrics reuse notice" for n in ns)
        dec = sorted(day(n["notice_date"]) for n in ns if n["notice_type"] in ("Approval notice", "Denial notice"))
        wd = sorted(day(n["notice_date"]) for n in ns if n["notice_type"] == "Withdrawal acknowledgement")
        decided = dec[0] if dec else None
        if wd and (decided is None or wd[0] < decided):
            continue
        c = case.get(k, {})
        st = stop.get(k) or (c.get("closed") if c.get("status") == "Closed - no longer tracked" else None)
        rq = sorted(req.get(k, []))
        ren.append(dict(rcv=rcv, cat=rc[0]["eligibility_category"], appt=day(ap[-1]["appointment_date"]) if ap else None,
                        reuse=reuse and not ap, decided=decided, stop=st, card=c.get("card"), req=rq[-1] if rq else None))
    ren = [r for r in ren if r["appt"] or r["reuse"]]
    pA = {cat: np.mean([r["appt"] is not None for r in ren if r["cat"] == cat]) for cat in ("(c)(9)", "(c)(26)")}
    L = np.array([(r["appt"] - r["rcv"]).days for r in ren if r["appt"] and LEAD_WINDOW[0] <= r["rcv"] <= LEAD_WINDOW[1]])
    qt, qe = [], []
    for r in ren:
        s = r["appt"] if r["appt"] else r["rcv"]
        if s > EXPORT:
            continue
        end = min(x for x in (EXPORT, r["stop"]) if x)
        if r["req"] and (r["decided"] is None or r["req"][0] < r["decided"]):
            t, ev = (r["req"][0] - s).days, 0
        elif r["decided"] and r["decided"] <= end:
            t, ev = (r["decided"] - s).days, 1
        else:
            t, ev = (end - s).days, 0
        if t >= 0:
            qt.append(t)
            qe.append(ev)
    SQ = kaplan_meier(qt, qe)
    post = [r for r in ren if r["rcv"] >= RULE]
    gap = np.array([(r["card"] - r["rcv"]).days for r in post if r["card"]])
    with_req = [r for r in post if r["req"]]
    outs = [r["req"][1] for r in with_req if r["req"][1] in ("Accepted", "Denied")]
    acc = outs.count("Accepted") / len(outs)
    E = np.array([(r["decided"] - r["req"][0]).days for r in with_req if r["req"][1] == "Accepted" and r["decided"] and
                  r["decided"] >= r["req"][0]])
    OFF = np.array([(r["req"][0] - r["card"]).days for r in with_req if r["card"]])

    # ------------------------------------------------------------------- ask A's records: each current person's permit
    emp = rows_csv(os.path.join(target, "employees_2026-12.csv"))
    dep = rows_csv(os.path.join(target, "dependants_2026-12.csv"))
    permits = rows_xlsx(os.path.join(target, "work_permits_2026-12.xlsx"), "Permits")
    tracker = rows_xlsx(os.path.join(target, "green_card_tracker_2026-12.xlsx"), "Applicants")
    active = {e["employee_id"] for e in emp if e["status"] == "Active"}
    employees = {e["employee_id"] for e in emp}
    sponsor = {}
    for p in permits:
        if p.get("household_employee_id"):
            sponsor.setdefault(p["holder_id"], p["household_employee_id"])
    for c in case.values():
        if c["ptype"] == "Spouse" and c["sponsor"]:
            sponsor.setdefault(c["person"], c["sponsor"])
    for t in tracker:
        if t["relationship"] == "Spouse":
            sponsor.setdefault(t["applicant_id"], t["principal_employee_id"])
    enr = {e["benefits_enrollment_id"]: e["employee_id"] for e in emp}
    for d in dep:
        if d["relationship"] != "Child" and d["benefits_enrollment_id"] in enr:
            sponsor.setdefault(d["dependant_id"], enr[d["benefits_enrollment_id"]])
    cards = {}                      # (holder, first day) -> [last day, category, from a filing]

    def put(h, vf, vt, cat, filing):
        if not (h and vf and vt) or vf > EXPORT:
            return
        cur = cards.get((h, vf))
        if cur is None:
            cards[(h, vf)] = [vt, cat, filing]
        elif filing and not cur[2]:
            cards[(h, vf)] = [vt, cat, True]
    for p in permits:
        put(p["holder_id"], day(p["valid_from"]), day(p["valid_to"]), p["category"], False)
    for c in case.values():
        if c["status"] == "Approved" and c["new"] and c["decided"]:
            put(c["person"], c["decided"], c["new"], c["cat"], True)
    for n in notices:
        if n["notice_type"] == "Approval notice":
            c = case.get(key.get(n["client_ref"]))
            if c:
                put(c["person"], day(n["valid_from"]), day(n["valid_to"]), n["eligibility_category"], True)
    for t in tracker:
        if t["ead_approved"] and t["ead_valid_to"]:
            put(t["applicant_id"], day(t["ead_approved"]), day(t["ead_valid_to"]), "(c)(9)", True)
    permit = {}                     # person -> the card approved most recently; a spouse is her household (an employee
    for (h, vf), (vt, cat, _) in cards.items():     # has one), whichever ID a record files her card under
        person = h if h in employees else ("spouse of", sponsor.get(h))
        cur = (h in active) if h in employees else (sponsor.get(h) in active)
        if not cur:
            continue
        if person not in permit or vf > permit[person][0]:
            permit[person] = (vf, vt, cat, _)
        elif vf == permit[person][0] and _ and not permit[person][3]:
            permit[person] = (vf, vt, cat, True)     # one card under two IDs: the approval's category stands
    permit = {k: v[:3] for k, v in permit.items()}

    # the quarter's categories: current permits whose renewal reaches USCIS in the quarter (last day less the median
    # card gap), less the (c)(9) permits of applicants whose adjustment is approved
    lpr = {t["applicant_id"] if t["relationship"] == "Principal" else ("spouse of", t["principal_employee_id"])
           for t in tracker if t["i485_status"] == "Approved"}
    g50 = int(round(float(np.median(gap))))
    nmix = defaultdict(int)
    for h, (vf, vt, cat) in permit.items():
        if cat == "(c)(9)" and h in lpr:
            continue
        if QUARTER[0] <= vt - dt.timedelta(days=g50) <= QUARTER[1]:
            nmix[cat] += 1
    tot = sum(nmix.values())
    mix = {cat: nmix[cat] / tot for cat in ("(c)(9)", "(c)(26)")}
    pA_mix = sum(mix[c] * pA[c] for c in mix)

    H = 900

    def sq(k):
        k = np.asarray(k)
        out = np.ones(k.shape)
        pos = k >= 0
        out[pos] = np.where(k[pos] < len(SQ), SQ[np.minimum(k[pos], len(SQ) - 1)], SQ[-1])
        return out
    w = np.arange(H)
    pl = np.bincount(L.astype(int), minlength=H)[:H] / len(L)
    s_appt = np.array([np.sum(pl * sq(x - np.arange(H))) for x in w])
    s_nat = (1 - pA_mix) * sq(w) + pA_mix * s_appt
    lo = int(min(0, OFF.min()))
    pd_ = np.bincount(gap.astype(int), minlength=H)[:H] / len(gap)
    po = np.bincount((OFF - lo).astype(int)) / len(OFF)
    pe = np.bincount(E.astype(int), minlength=H)[:H] / len(E)
    conv = np.convolve(np.convolve(pd_, po), pe)
    g = np.cumsum(conv)                           # index i of the convolution is day i + lo
    G = np.array([g[min(i - lo, len(g) - 1)] if i - lo >= 0 else 0.0 for i in range(H)])
    S_w = s_nat * (1 - acc * G)
    q80 = int(np.argmax(1 - S_w >= 0.8))
    months = q80 / MONTH
    quote = np.floor(months * 2 + 0.5) / 2
    say(quote == figs["quote_months"], "the quote recomputes by exact convolution", "%d days, %.3f months, %.1f (the "
        "quarter %.0f%% (c)(9), appointment share %.3f)" % (q80, months, quote, 100 * mix["(c)(9)"], pA_mix))
    edge = 0.25 - abs(months - quote)
    say(edge >= 0.125, "the quote sits a quarter-bin clear of its edges", "edge distance %.3f" % edge)
    say(round(np.median(L) / 7) == figs["lead_weeks_now"], "the appointment lead now set", "%.2f weeks" % (
        np.median(L) / 7))
    oldL = [(r["appt"] - r["rcv"]).days for r in ren if r["appt"] and r["rcv"] < dt.date(2026, 1, 1)]
    say(round(np.median(oldL) / 7) == figs["lead_weeks_before"], "the lead through 2024 and 2025", "%.2f weeks" % (
        np.median(oldL) / 7))
    # appointment-route renewals of the quarter past their card's expiry undecided: P(L + Q > D)
    pdist = np.bincount(gap.astype(int), minlength=H)[:H] / len(gap)
    reach_appt = float(np.sum(pdist * s_appt))
    say(abs(reach_appt - figs["reach_appt_share"]) < 0.03 and 0.85 <= reach_appt < 0.95,
        "about nine in ten appointment-route renewals reach the card's expiry undecided", "%.3f" % reach_appt)
    say(round(np.median(E) / 7) == figs["expedite_response_weeks"], "the expedite response, about a week",
        "%.1f days" % np.median(E))
    # the rival the write-up names: the record's own 80th percentile, open renewals censored
    tt, ee = [], []
    for r in ren:
        end = min(x for x in (EXPORT, r["stop"]) if x)
        if r["decided"] and r["decided"] <= end:
            tt.append((r["decided"] - r["rcv"]).days)
            ee.append(1)
        else:
            tt.append((end - r["rcv"]).days)
            ee.append(0)
    S_rec = kaplan_meier(tt, ee)
    rec80 = int(np.argmax(S_rec <= 0.2 + 1e-12))
    say(np.floor(rec80 / MONTH * 2 + 0.5) / 2 == figs["rival_record_months"], "the record's own 80th percentile",
        "%.3f months" % (rec80 / MONTH))

    # ------------------------------------------------------------------- ask A
    got = defaultdict(int)
    for h, (vf, vt, cat) in permit.items():
        if vt.year == 2027 and vt.month in (7, 8, 9):
            got["2027-%02d %s" % (vt.month, cat)] += 1
    for k, v in figs["ask_a"].items():
        say(got[k] == v, "ask A " + k, "%d" % got[k])

    # ------------------------------------------------------------------- ask B
    rej = set()
    for n in notices:
        if n["notice_type"] == "Rejection notice":
            rej.add((key.get(n["client_ref"]), day(n["notice_date"])))
    for a in actions:
        if a["action"] == "Notice forwarded by applicant" and a["detail"] == "Rejection notice":
            rej.add((key.get(a["case_id"]), day(a["notice_date"])))
    rcv = defaultdict(set)
    for n in notices:
        if n["notice_type"] == "Receipt notice":
            rcv[key.get(n["client_ref"])].add(day(n["received_date"]))
    for k, c in case.items():
        if c["kind"] == "renewal" and c["received"]:
            rcv[k].add(c["received"])
    quarters = {"2025 Q4": (dt.date(2025, 10, 1), dt.date(2025, 12, 31)), "2026 Q1": (dt.date(2026, 1, 1), dt.date(2026, 3, 31)),
                "2026 Q2": (dt.date(2026, 4, 1), dt.date(2026, 6, 30)), "2026 Q3": (dt.date(2026, 7, 1), dt.date(2026, 9, 30))}
    for lab, (a, b) in quarters.items():
        inq = [(k, d) for k, d in rej if a <= d <= b]
        ref30 = 0
        for k, d in inq:
            nxt = sorted(x for x in rcv.get(k, ()) if x > d)
            ref30 += bool(nxt and (nxt[0] - d).days <= 30)
        say(len(inq) == figs["ask_b"][lab]["rejected"], "ask B %s rejected" % lab, "%d" % len(inq))
        say(ref30 == figs["ask_b"][lab]["refiled_30"], "ask B %s filed again within 30 days" % lab, "%d" % ref30)

    # ------------------------------------------------------------------- the goldens carry the same figures
    gold = os.path.join(root, "golden")
    ws = load_workbook(os.path.join(gold, "q1_2027_renewal_quote.xlsx"), read_only=True)
    q_cell = [r for r in ws["Quote"].iter_rows(values_only=True) if r and r[0] == "Wait to quote"][0][1]
    say(q_cell == "%.1f months" % figs["quote_months"], "the workbook states the quote", q_cell)
    exp_rows = [r for r in ws["Expiring permits"].iter_rows(values_only=True) if r and isinstance(r[1], int) and
                not str(r[0]).startswith("July to")]
    say([(r[1], r[2]) for r in exp_rows] == [(figs["ask_a"]["2027-%02d (c)(9)" % m], figs["ask_a"]["2027-%02d (c)(26)" % m])
                                            for m in (7, 8, 9)], "the workbook's expiring permits")
    rej_rows = [r for r in ws["Lockbox rejections"].iter_rows(values_only=True) if r and isinstance(r[1], int) and
                not str(r[0]).startswith("October 2025 to")]
    say([(r[1], r[2]) for r in rej_rows] == [(figs["ask_b"][q]["rejected"], figs["ask_b"][q]["refiled_30"])
                                            for q in quarters], "the workbook's lockbox rejections")
    svg = open(os.path.join(gold, "q1_2027_expiring_permits.svg"), encoding="utf-8").read()
    say(all((">%d<" % v) in svg for v in figs["ask_a"].values()), "the chart labels every bar with its count")
    sub = open(os.path.join(root, "submission.md"), encoding="utf-8").read()
    say("%.1f months" % figs["quote_months"] in sub, "the write-up states the quote at the half month")
    print("%d failed" % len(FAILS))
    return 1 if FAILS else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
