"""task172 pack: every file under target/, written from the records (records.py) the way each source system exports
them, then made reproducible (postprocess.py)."""
from __future__ import annotations

import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import common as C          # noqa: E402
import docs as D            # noqa: E402
import postprocess as PP    # noqa: E402
import records as Rm        # noqa: E402
import world as Wm          # noqa: E402

HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
STAMPS = {
    C.F_REGISTER: ("Melinda Yang", "Brandon Zamora", "2023-11-27T17:14:00Z", "2026-12-04T22:41:00Z"),
    C.F_PERMITS: ("HR Systems", "HR Systems", "2021-02-01T15:03:00Z", "2026-12-04T13:20:00Z"),
    C.F_TRACKER: ("Brandon Zamora", "Brandon Zamora", "2021-01-11T18:25:00Z", "2026-12-03T23:52:00Z"),
    C.F_H1B: ("Raymond Navarro", "Brandon Zamora", "2026-08-03T16:40:00Z", "2026-11-20T18:06:00Z"),
    C.F_DICT: ("Brandon Zamora", "Melinda Yang", "2024-02-12T19:30:00Z", "2026-12-04T23:08:00Z"),
}


def _col(j):
    s = ""
    j += 1
    while j:
        j, r = divmod(j - 1, 26)
        s = chr(65 + r) + s
    return s


def _sheet(ws, head, rows, widths=None, date_fmt="yyyy-mm-dd", bold=True):
    for j, h in enumerate(head):
        c = ws.cell(row=1, column=j + 1, value=h)
        if bold:
            c.font = Font(bold=True)
            c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="top")
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            if v is None or v == "":
                continue
            c = ws.cell(row=i + 2, column=j + 1, value=v)
            if isinstance(v, dt.date):
                c.number_format = date_fmt
    if widths:
        for j, w in enumerate(widths):
            ws.column_dimensions[_col(j)].width = w
    ws.freeze_panes = ws.cell(row=2, column=1)


def _save(wb, path, key):
    wb.save(path)
    creator, last_by, created, modified = STAMPS[key]
    PP.normalize_ooxml(path, creator, last_by, created, modified, excel=True, company=C.ORG)


def write_notices(out, R):
    rows = []
    for n in R.notices:
        cs = n["cs"]
        rows.append((n["id"], n["matter"], cs.case_ref, n["receipt"] or "", n["form"], cs.cat, n["ntype"],
                     C.d_iso(n["ndate"]), C.d_iso(n["received"]), C.d_iso(n["appt"]), n["site"] if n["appt"] else "",
                     C.d_iso(n["vf"]), C.d_iso(n["vt"]), n["office"], n["src"], C.d_iso(n["logged"])))
    return C.write_csv(os.path.join(out, C.F_NOTICES),
                       ("notice_id", "matter_no", "client_ref", "receipt_no", "form", "eligibility_category",
                        "notice_type", "notice_date", "received_date", "appointment_date", "appointment_site",
                        "valid_from", "valid_to", "uscis_office", "source", "logged_on"), rows)


def write_register(out, R):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cases to Mar 2025"
    old = []
    for r in R.reg_old:
        old.append((r["id"], "Renewal" if r["cs"] is not None else "Initial", r["person"], r["ptype"], r["sponsor"],
                    r["cat"], C.d_us(r["card_exp"]), C.d_us(r["opened"]), C.d_us(r["sent"]), r["matter"], r["receipt"],
                    C.d_us(r["received"]), r["status"], C.d_us(r["decided"]), C.d_us(r["new_exp"]), C.d_us(r["closed"]),
                    r["filed_by"]))
    _sheet(ws, ("Case ID", "Case Type", "Person ID", "Person Type", "Sponsor ID", "Category", "Card Expires", "Opened",
                "Packet Sent", "Matter No.", "Receipt No.", "Received", "Status", "Decided", "New Card Expires",
                "Closed", "Filed By"), old,
           widths=(13, 10, 11, 11, 11, 9, 12, 11, 11, 15, 15, 11, 22, 11, 15, 11, 19))
    ws2 = wb.create_sheet("Cases")
    new = []
    for r in R.reg_new:
        new.append((r["id"], r["legacy"] or None, "renewal" if r["cs"] is not None else "initial", r["person"],
                    r["ptype"].lower(), r["sponsor"] or None, r["cat"], r["card_exp"], r["opened"], r["sent"],
                    r["matter"] or None, r["receipt"] or None, r["received"], r["status"], r["decided"], r["new_exp"],
                    r["closed"], r["filed_by"]))
    _sheet(ws2, ("case_id", "legacy_case_id", "case_type", "person_id", "person_type", "sponsor_id", "category",
                 "card_expires", "opened_on", "packet_sent_on", "matter_no", "receipt_no", "received_on", "status",
                 "decided_on", "new_card_expires", "closed_on", "filed_by"), new,
           widths=(11, 13, 9, 11, 10, 11, 9, 12, 11, 13, 15, 15, 12, 24, 11, 15, 11, 19))
    _save(wb, os.path.join(out, C.F_REGISTER), C.F_REGISTER)
    return len(old) + len(new)


def write_actions(out, R):
    rows = [(a["id"], a["case"], a["person"], C.d_iso(a["when"]), a["action"], a["detail"], a["ref"], a["outcome"],
             C.d_iso(a["out_date"]), C.d_iso(a["notice_date"]), a["by"]) for a in R.actions]
    return C.write_csv(os.path.join(out, C.F_ACTIONS),
                       ("action_id", "case_id", "person_id", "action_date", "action", "detail", "reference", "outcome",
                        "outcome_date", "notice_date", "entered_by"), rows)


def write_permits(out, R):
    wb = Workbook()
    ws = wb.active
    ws.title = "Permits"
    rows = [(r["holder"], r["htype"], r["cat"], r["number"], r["vf"], r["vt"], r["ent"], r["via"], r["sponsor"] or None)
            for r in R.roster]
    _sheet(ws, ("holder_id", "holder_type", "category", "card_number", "valid_from", "valid_to", "entered_on",
                "entered_via", "household_employee_id"), rows, widths=(11, 11, 9, 16, 11, 11, 11, 19, 14))
    ws.auto_filter.ref = "A1:I%d" % (len(rows) + 1)
    _save(wb, os.path.join(out, C.F_PERMITS), C.F_PERMITS)
    return len(rows)


def write_employees(out, R):
    rows = [(e["id"], C.d_iso(e["hire"]), e["status"], C.d_iso(e["term"]), e["state"], e["family"], e["basis"],
             e["enroll"]) for e in R.employees]
    return C.write_csv(os.path.join(out, C.F_EMPLOYEES),
                       ("employee_id", "hire_date", "status", "termination_date", "work_state", "job_family",
                        "immigration_basis", "benefits_enrollment_id"), rows)


def write_dependants(out, R):
    rows = [(d["id"], d["enroll"], d["rel"], C.d_iso(d["start"]), d["birth"]) for d in R.dependants]
    return C.write_csv(os.path.join(out, C.F_DEPENDANTS),
                       ("dependant_id", "benefits_enrollment_id", "relationship", "coverage_start", "birth_year"), rows)


def write_tracker(out, R):
    wb = Workbook()
    ws = wb.active
    ws.title = "Applicants"
    rows = [(t["applicant"], t["principal"], t["rel"], t["country"], t["pd"], t["i140"], t["i485"], t["status"],
             t["decided"], t["ead_from"], t["ead_to"]) for t in R.tracker]
    _sheet(ws, ("applicant_id", "principal_employee_id", "relationship", "chargeability", "priority_date",
                "i140_approved", "i485_received", "i485_status", "i485_decided", "ead_approved", "ead_valid_to"), rows,
           widths=(12, 12, 11, 12, 12, 12, 12, 11, 12, 12, 12))
    _save(wb, os.path.join(out, C.F_TRACKER), C.F_TRACKER)
    return len(rows)


def write_dispatch(out, R):
    rows = [(r["pid"], C.d_iso(r["sent"]), r["case"], r["person"], r["contents"], r["to"], r["courier"],
             r["tracking"], r["status"], C.d_iso(r["returned"])) for r in R.dispatch]
    return C.write_csv(os.path.join(out, C.F_DISPATCH),
                       ("packet_id", "sent_on", "case_id", "person_id", "contents", "addressed_to", "courier",
                        "tracking_no", "status", "returned_on"), rows)


def write_posted(out, R):
    rows = [(C.d_iso(p["cap"]), p["line"], p["office"], "%.1f" % p["months"], C.d_iso(p["inquiry"])) for p in R.posted]
    return C.write_csv(os.path.join(out, C.F_POSTED),
                       ("captured_on", "form_line", "office", "posted_months", "case_inquiry_date"), rows)


def write_h1b(out, R):
    wb = Workbook()
    ws = wb.active
    ws.title = "2027"
    rows = [(r["id"], r["end"], r["plan"], r["kind"], r["h4"], r["premium"]) for r in R.h1b]
    _sheet(ws, ("Employee ID", "H-1B valid to", "Planned filing month", "Petition", "H-4 dependants",
                "Premium processing"), rows, widths=(12, 13, 19, 26, 14, 18))
    _save(wb, os.path.join(out, C.F_H1B), C.F_H1B)
    return len(rows)


def records():
    W = Wm.build()
    R = Rm.write_all(W)
    Rm.roster(W, R)
    Rm.people(W, R)
    Rm.h1b_calendar(W, R)
    Rm.posted_times(R)
    return W, R


def build_pack(out, WR=None):
    os.makedirs(out, exist_ok=True)
    W, R = WR or records()
    n = {}
    n[C.F_NOTICES] = write_notices(out, R)
    n[C.F_REGISTER] = write_register(out, R)
    n[C.F_ACTIONS] = write_actions(out, R)
    n[C.F_PERMITS] = write_permits(out, R)
    n[C.F_EMPLOYEES] = write_employees(out, R)
    n[C.F_DEPENDANTS] = write_dependants(out, R)
    n[C.F_TRACKER] = write_tracker(out, R)
    n[C.F_DISPATCH] = write_dispatch(out, R)
    n[C.F_POSTED] = write_posted(out, R)
    n[C.F_H1B] = write_h1b(out, R)
    D.write_policy(out)
    D.write_thread(out)
    n[C.F_DICT] = D.write_dictionary(out, STAMPS[C.F_DICT])
    n[C.F_MANIFEST] = D.write_manifest(out, n)
    return n


if __name__ == "__main__":
    print(build_pack(sys.argv[1]))
