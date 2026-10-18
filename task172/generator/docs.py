"""task172 documents: Global Mobility's renewal standard (PDF), the December planning thread (.eml), the data
dictionary (Excel) and the export manifest. Their text is held here as data so checks.py can quote it and verify.py
can read it back."""
from __future__ import annotations

import datetime as dt
import os

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                        # noqa: E402
from reportlab.lib.pagesizes import letter                                              # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                         # noqa: E402
from reportlab.lib.units import mm                                                      # noqa: E402
from reportlab.pdfbase import pdfmetrics                                                # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                            # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer                     # noqa: E402

import common as C          # noqa: E402
import postprocess as PP    # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/"
for _n, _f in (("Carlito", "crosextra/Carlito-Regular.ttf"), ("Carlito-Bold", "crosextra/Carlito-Bold.ttf"),
               ("LibSans", "liberation/LiberationSans-Regular.ttf"), ("LibSans-Bold", "liberation/LiberationSans-Bold.ttf")):
    pdfmetrics.registerFont(TTFont(_n, FONT_DIR + _f))
from reportlab.pdfbase.pdfmetrics import registerFontFamily   # noqa: E402
registerFontFamily("Carlito", normal="Carlito", bold="Carlito-Bold", italic="Carlito", boldItalic="Carlito-Bold")
INK = colors.HexColor("#1f2328")
GREY = colors.HexColor("#5d646d")
ACCENT = colors.HexColor("#0b5a6b")

POLICY_DATE = dt.date(2025, 11, 12)
POLICY_TITLE = "Work Permit Renewal Standard"
POLICY_SUB = "Global Mobility standard GM-07, revision 4, effective 12 November 2025. Owner: Jasmin Sims, Head of Global Mobility."
POLICY = [
    ("1. Scope", [
        "1.1 This standard covers the renewal of work permits (Form I-765) that Global Mobility files through outside "
        "counsel for employees whose adjustment of status is pending, in category (c)(9), and for the spouses of "
        "employees, in category (c)(26) while the employee holds H-1B status and in category (c)(9) while the spouse's "
        "own adjustment is pending.",
        "1.2 Global Mobility files renewals for employees and for the spouses of employees. When an employee leaves "
        "Sallowmere, work stops on the employee's and the spouse's pending renewals.",
        "1.3 Global Mobility tracks the work permits of current employees and of the spouses of current employees. An "
        "employee is current while HR's program roster shows the employee as Active. A spouse is the person Global "
        "Mobility took into its care with the employee's household, whether Benefits covers her as a spouse or as a "
        "domestic partner.",
        "1.4 A person's work permit is the card USCIS approved for the person most recently, in the category on its "
        "approval notice, which is not always the category entered on HR's work permit roster. It remains the "
        "person's work permit until its last day, whatever has changed in the person's status since.",
    ]),
    ("2. Opening and filing", [
        "2.1 Global Mobility opens a renewal 180 days before the current card expires and files it when the packet is "
        "complete.",
        "2.2 Packets go by courier to the USCIS lockbox for the filer's state of residence, with counsel's Form G-28.",
        "2.3 A renewal is filed on the date USCIS receives it, the receipt date on its receipt notice.",
        "2.4 A packet the lockbox rejects is not filed. A rejected renewal is filed again within 30 days when USCIS "
        "receives it no later than the 30th day after the date on the rejection notice.",
    ]),
    ("3. When a card expires", [
        "3.1 Renewals filed on or after 30 October 2025 do not extend the current card. An employee whose card expires "
        "before USCIS decides the renewal may not work from the day after it expires and is placed on immigration "
        "leave until the new card is in hand.",
        "3.2 The first ten business days of immigration leave are paid. Leave beyond that is unpaid, and benefits "
        "continue.",
    ]),
    ("4. The renewal wait in the benefits portal", [
        "4.1 Each quarter the global mobility analytics lead publishes in the benefits portal, from the 15th of the "
        "last month of the quarter before, the wait within which four in five of the renewals Global Mobility files "
        "in the coming quarter will have USCIS's decision.",
        "4.2 The wait runs from the date a renewal is filed to the date of USCIS's decision notice, approval or "
        "denial. Renewals withdrawn before a decision are left out.",
        "4.3 Waits are quoted in months of 30.4 days, to the nearest half month.",
    ]),
    ("5. Review", [
        "5.1 Global Mobility reviews this standard each year and when USCIS changes how renewals are filed or "
        "extended.",
    ]),
]
REVISIONS = ("Revision 4 (12 November 2025): section 3 rewritten for the end of automatic extension. "
             "Revision 3 (2 April 2024): section 2.2, the lockbox by state of residence.")


def section_numbers():
    return [p.split(" ", 1)[0] for _, paras in POLICY for p in paras if p[:1].isdigit()]


def policy_text():
    out = [POLICY_TITLE, C.ORG, POLICY_SUB]
    for h, paras in POLICY:
        out.append(h)
        out.extend(paras)
    out.append(REVISIONS)
    return out


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def pdf_stamp(d, hh=10, mi=0):
    off = -8 if d.month in (1, 2, 11, 12) or (d.month == 3 and d.day < 9) else -7
    return ("D:%04d%02d%02d%02d%02d00%+03d'00'" % (d.year, d.month, d.day, hh, mi, off)).encode()


def _footer(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Carlito", 7.8)
        canvas.setFillColor(GREY)
        canvas.drawString(22 * mm, 11 * mm, text)
        canvas.drawRightString(letter[0] - 22 * mm, 11 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def write_policy(out):
    path = os.path.join(out, C.F_POLICY)
    st = dict(
        org=ParagraphStyle("org", fontName="Carlito-Bold", fontSize=9.5, leading=12, textColor=ACCENT, spaceAfter=2),
        h1=ParagraphStyle("h1", fontName="LibSans-Bold", fontSize=16, leading=20, textColor=INK, spaceAfter=4),
        sub=ParagraphStyle("sub", fontName="Carlito", fontSize=9.2, leading=12, textColor=GREY, spaceAfter=12),
        h2=ParagraphStyle("h2", fontName="LibSans-Bold", fontSize=11, leading=14, textColor=ACCENT, spaceBefore=8,
                          spaceAfter=3),
        body=ParagraphStyle("body", fontName="Carlito", fontSize=11, leading=14.6, textColor=INK, spaceAfter=5),
        small=ParagraphStyle("small", fontName="Carlito", fontSize=8.6, leading=11, textColor=GREY, spaceBefore=10))
    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=22 * mm, rightMargin=22 * mm, topMargin=20 * mm,
                            bottomMargin=20 * mm, title=POLICY_TITLE, author="Global Mobility", subject="", creator="",
                            producer="")
    story = [Paragraph(esc(C.ORG + " | " + C.TEAM), st["org"]), Paragraph(esc(POLICY_TITLE), st["h1"]),
             Paragraph(esc(POLICY_SUB), st["sub"])]
    for h, paras in POLICY:
        story.append(Paragraph(esc(h), st["h2"]))
        for p in paras:
            story.append(Paragraph(esc(p), st["body"]))
    story.append(Spacer(1, 4))
    story.append(Paragraph(esc(REVISIONS), st["small"]))
    foot = "Sallowmere Software internal. GM-07 rev. 4"
    doc.build(story, onFirstPage=_footer(foot), onLaterPages=_footer(foot))
    PP.normalize_pdf(path, pdf_stamp(POLICY_DATE, 11, 4), pdf_stamp(POLICY_DATE, 11, 4))


# ------------------------------------------------------------------------------------------------------ the thread
THREAD = [
    dict(frm="Brandon Zamora <bzamora@sallowmere.com>", to="Melinda Yang <myang@sallowmere.com>",
         cc="Jasmin Sims <jsims@sallowmere.com>", date="Thu, 03 Dec 2026 16:42:19 -0800",
         body=["Melinda,",
               "",
               "The export from Marrowby's system runs tomorrow morning. I'll put it in the Mobility share with our "
               "register, the case log, the roster pull from HR and the mailroom log.",
               "",
               "Appointment letters for the October and November filings are mostly for late January into March. The "
               "(c)(9)s come back in January too, the first renewals of the 18-month cards.",
               "",
               "Brandon"]),
    dict(frm="Melinda Yang <myang@sallowmere.com>", to="Brandon Zamora <bzamora@sallowmere.com>",
         cc="Jasmin Sims <jsims@sallowmere.com>", date="Wed, 02 Dec 2026 09:15:52 -0800",
         body=["Brandon,",
               "",
               "The 15 December release closes for edits on the 10th. For the renewals page I need the Q1 wait, the "
               "expiring-permits chart we ran last time (July to September this time) and the lockbox numbers Jasmin "
               "asked for at the quarterly review.",
               "",
               "Can you send me what you have?",
               "",
               "Thanks,",
               "Melinda"]),
]


def thread_text():
    return [ln for m in THREAD for ln in m["body"] if ln]


def write_thread(out):
    last = THREAD[0]
    lines = ["Return-Path: <bzamora@sallowmere.com>",
             "Message-ID: <CAF4q2v7Hk0aZ1@mail.sallowmere.com>",
             "In-Reply-To: <CAB8m1s3Rt9wQ4@mail.sallowmere.com>",
             "References: <CAB8m1s3Rt9wQ4@mail.sallowmere.com>",
             "Date: " + last["date"], "From: " + last["frm"], "To: " + last["to"], "Cc: " + last["cc"],
             "Subject: RE: Q1 renewal wait for the portal", "MIME-Version: 1.0",
             "Content-Type: text/plain; charset=\"utf-8\"", "Content-Transfer-Encoding: 8bit", ""]
    lines.extend(last["body"])
    prev = THREAD[1]
    lines.extend(["", "________________________________", "From: Melinda Yang", "Sent: Wednesday, December 2, 2026 9:15 AM",
                  "To: Brandon Zamora",
                  "Cc: Jasmin Sims", "Subject: Q1 renewal wait for the portal", ""])
    lines.extend("> " + ln if ln else ">" for ln in prev["body"])
    with open(os.path.join(out, C.F_THREAD), "w", encoding="utf-8", newline="") as fh:
        fh.write("\r\n".join(lines) + "\r\n")


# ------------------------------------------------------------------------------------------------------ dictionary
DICT = [
    (C.F_NOTICES, "(file)", "Counsel's export of the USCIS notices and case-status events logged on Global Mobility's "
     "renewal matters. One row per notice or event as logged. A matter logs nothing after counsel closes it."),
    (C.F_NOTICES, "notice_id", "Counsel's identifier for the logged notice or event."),
    (C.F_NOTICES, "matter_no", "Counsel's matter number."),
    (C.F_NOTICES, "client_ref", "The Global Mobility case reference given to counsel when the matter was opened."),
    (C.F_NOTICES, "receipt_no", "USCIS receipt number of the filing."),
    (C.F_NOTICES, "form", "Forms the notice covers."),
    (C.F_NOTICES, "eligibility_category", "Work permit category applied for."),
    (C.F_NOTICES, "notice_type", "The notice or status event (sheet Codes)."),
    (C.F_NOTICES, "notice_date", "Date on the notice, or the date of the status event."),
    (C.F_NOTICES, "received_date", "On a receipt notice, the date USCIS received the filing. On a rejection notice, "
     "the date the lockbox received the packet."),
    (C.F_NOTICES, "appointment_date", "On a biometrics appointment notice, the date of the appointment."),
    (C.F_NOTICES, "appointment_site", "Application support center of the appointment."),
    (C.F_NOTICES, "valid_from", "On an approval notice, the first day of the new card."),
    (C.F_NOTICES, "valid_to", "On an approval notice, the last day of the new card."),
    (C.F_NOTICES, "uscis_office", "USCIS office or lockbox that issued the notice."),
    (C.F_NOTICES, "source", "How counsel received it: mail scan, the USCIS online account, or the case status feed."),
    (C.F_NOTICES, "logged_on", "Date counsel logged the notice."),
    (C.F_REGISTER, "(file)", "Global Mobility's case register. Cases opened before 1 April 2025 are in sheet Cases to "
     "Mar 2025 as that system held them on 31 March 2025, with dates as text, month first. Cases open on that day "
     "were carried into the current system (sheet Cases) under a GM reference, with the RN reference in "
     "legacy_case_id."),
    (C.F_REGISTER, "case_id", "Global Mobility case reference (Case ID in the earlier sheet)."),
    (C.F_REGISTER, "case_type", "Initial or renewal filing."),
    (C.F_REGISTER, "person_id", "Employee ID; for a spouse, the benefits dependant ID, or a Global Mobility person "
     "ID (MP-) where the spouse was not on Sallowmere benefits when the household came into Global Mobility's care."),
    (C.F_REGISTER, "sponsor_id", "For a spouse, the employee whose household the case belongs to."),
    (C.F_REGISTER, "category", "Work permit category of the filing."),
    (C.F_REGISTER, "card_expires", "On a renewal, the expiry date of the card being renewed."),
    (C.F_REGISTER, "opened_on", "Date Global Mobility opened the case."),
    (C.F_REGISTER, "packet_sent_on", "Date the first packet left the mailroom."),
    (C.F_REGISTER, "matter_no", "Counsel's matter number."),
    (C.F_REGISTER, "receipt_no", "USCIS receipt number, entered from the receipt notice."),
    (C.F_REGISTER, "received_on", "USCIS receipt date, entered from the receipt notice."),
    (C.F_REGISTER, "status", "Open, Sent, Filed, Approved, Denied, Withdrawn, Closed - no longer tracked, Rejected - "
     "not refiled; Migrated in the earlier sheet for a case carried into the current system."),
    (C.F_REGISTER, "decided_on", "Date of USCIS's decision, entered from the decision notice."),
    (C.F_REGISTER, "new_card_expires", "Expiry of the card issued on approval."),
    (C.F_REGISTER, "closed_on", "Date the case closed."),
    (C.F_REGISTER, "filed_by", "Marrowby Teale, or Applicant where the filer sent the filing without counsel."),
    (C.F_ACTIONS, "(file)", "Global Mobility's log of its own actions on cases. One row per action."),
    (C.F_ACTIONS, "case_id", "Case reference in use on the action date."),
    (C.F_ACTIONS, "person_id", "Person the case is for."),
    (C.F_ACTIONS, "action_date", "Date of the action."),
    (C.F_ACTIONS, "action", "Action taken (sheet Codes)."),
    (C.F_ACTIONS, "detail", "Coordinator's note."),
    (C.F_ACTIONS, "reference", "USCIS service request number, or AR-11 confirmation number."),
    (C.F_ACTIONS, "outcome", "Outcome where the action has one."),
    (C.F_ACTIONS, "outcome_date", "Date of the outcome."),
    (C.F_ACTIONS, "notice_date", "For a notice forwarded by an applicant, the date on the notice."),
    (C.F_ACTIONS, "entered_by", "Coordinator's initials."),
    (C.F_PERMITS, "(file)", "HR's work permit roster. One row per card entered; rows stay when a card expires or a "
     "holder leaves."),
    (C.F_PERMITS, "holder_id", "Employee ID, benefits dependant ID, or Global Mobility person ID (MP-)."),
    (C.F_PERMITS, "holder_type", "Employee or Spouse."),
    (C.F_PERMITS, "category", "Category recorded for the card."),
    (C.F_PERMITS, "card_number", "Card number."),
    (C.F_PERMITS, "valid_from", "First day of the card."),
    (C.F_PERMITS, "valid_to", "Last day of the card."),
    (C.F_PERMITS, "entered_on", "Date the card was entered."),
    (C.F_PERMITS, "entered_via", "I-9 at hire; I-9 reverification; Onboarding, a spouse's card entered when Global "
     "Mobility takes the household into its care; Portal upload, a card a spouse uploads to the benefits portal."),
    (C.F_PERMITS, "household_employee_id", "On an onboarding entry, the employee whose household the spouse belongs "
     "to."),
    (C.F_EMPLOYEES, "(file)", "Employees in Global Mobility's programs, active on 4 December 2026 or terminated since "
     "1 January 2024."),
    (C.F_EMPLOYEES, "employee_id", "Employee ID."),
    (C.F_EMPLOYEES, "hire_date", "Most recent hire date."),
    (C.F_EMPLOYEES, "status", "Active or Terminated on 4 December 2026."),
    (C.F_EMPLOYEES, "termination_date", "Last day of employment."),
    (C.F_EMPLOYEES, "work_state", "State of the work location."),
    (C.F_EMPLOYEES, "job_family", "Job family."),
    (C.F_EMPLOYEES, "immigration_basis", "Current basis of the employee's authorization to work."),
    (C.F_EMPLOYEES, "benefits_enrollment_id", "The employee's benefits enrollment."),
    (C.F_DEPENDANTS, "(file)", "Dependants covered on Sallowmere benefits on 4 December 2026, continuation coverage "
     "included."),
    (C.F_DEPENDANTS, "dependant_id", "Benefits dependant ID."),
    (C.F_DEPENDANTS, "benefits_enrollment_id", "Enrollment the dependant is covered under."),
    (C.F_DEPENDANTS, "relationship", "Spouse, Domestic partner or Child."),
    (C.F_DEPENDANTS, "coverage_start", "First day of coverage."),
    (C.F_DEPENDANTS, "birth_year", "Year of birth."),
    (C.F_TRACKER, "(file)", "Global Mobility's green-card tracker. One row per applicant for adjustment of status, the "
     "principal and each derivative."),
    (C.F_TRACKER, "applicant_id", "Employee ID, benefits dependant ID, or Global Mobility person ID (MP-)."),
    (C.F_TRACKER, "principal_employee_id", "The employee whose petition the application rests on."),
    (C.F_TRACKER, "relationship", "Principal or Spouse."),
    (C.F_TRACKER, "chargeability", "Country of chargeability."),
    (C.F_TRACKER, "priority_date", "Priority date."),
    (C.F_TRACKER, "i140_approved", "Approval of the immigrant petition."),
    (C.F_TRACKER, "i485_received", "USCIS receipt date of the adjustment application."),
    (C.F_TRACKER, "i485_status", "Pending, Approved, Denied or Abandoned."),
    (C.F_TRACKER, "i485_decided", "Date of the decision on the adjustment application."),
    (C.F_TRACKER, "ead_approved", "Approval date of the applicant's latest (c)(9) work permit."),
    (C.F_TRACKER, "ead_valid_to", "Last day of that work permit."),
    (C.F_DISPATCH, "(file)", "The mailroom's dispatch log for 2025 and 2026, from the dispatch system used to 1 March "
     "2026 and the current system from 2 March 2026."),
    (C.F_DISPATCH, "packet_id", "Packet number. Each system numbers packets from 1, the earlier one from 1 each "
     "January."),
    (C.F_DISPATCH, "sent_on", "Date the packet left the mailroom."),
    (C.F_DISPATCH, "case_id", "Case reference the packet was sent under."),
    (C.F_DISPATCH, "person_id", "Person the packet was sent for."),
    (C.F_DISPATCH, "contents", "Forms in the packet."),
    (C.F_DISPATCH, "addressed_to", "Lockbox the packet was sent to."),
    (C.F_DISPATCH, "courier", "Courier."),
    (C.F_DISPATCH, "tracking_no", "Courier tracking number."),
    (C.F_DISPATCH, "status", "Sent; Receipted, a USCIS receipt notice is in hand; Returned, the packet came back to the "
     "mailroom."),
    (C.F_DISPATCH, "returned_on", "Date the packet came back to the mailroom."),
    (C.F_POSTED, "(file)", "Global Mobility's monthly capture of the processing time USCIS posts for Form I-765. One "
     "row per form line and office captured; office by office from July 2025."),
    (C.F_POSTED, "captured_on", "Date of the capture."),
    (C.F_POSTED, "form_line", "Form line as posted."),
    (C.F_POSTED, "office", "USCIS office."),
    (C.F_POSTED, "posted_months", "Time posted for the form line and office, in months."),
    (C.F_POSTED, "case_inquiry_date", "Receipt date posted for a case inquiry; captured to June 2025."),
    (C.F_H1B, "(file)", "Immigration team's 2027 H-1B extension calendar: employees whose H-1B status ends in 2027."),
    (C.F_H1B, "Employee ID", "Employee ID."),
    (C.F_H1B, "H-1B valid to", "Last day of the current H-1B approval."),
    (C.F_H1B, "Planned filing month", "Month the extension is planned to be filed."),
    (C.F_H1B, "Petition", "Petition type planned."),
    (C.F_H1B, "H-4 dependants", "Dependants in H-4 status."),
    (C.F_H1B, "Premium processing", "Whether premium processing is planned."),
]
CODES = [
    ("notice_type", "Receipt notice", "USCIS has received the filing."),
    ("notice_type", "Biometrics appointment notice", "An appointment at an application support center; a later "
     "notice replaces an earlier one."),
    ("notice_type", "Biometrics reuse notice", "USCIS will reuse biometrics it already holds."),
    ("notice_type", "Fingerprints taken", "Biometrics were taken."),
    ("notice_type", "Request for evidence", "USCIS asked for more evidence."),
    ("notice_type", "Evidence received", "USCIS received the response to its request."),
    ("notice_type", "Case transferred", "The case moved to another USCIS office."),
    ("notice_type", "Case under review", "An officer is reviewing the case."),
    ("notice_type", "Approval notice", "The filing was approved."),
    ("notice_type", "Denial notice", "The filing was denied."),
    ("notice_type", "Withdrawal acknowledgement", "USCIS acknowledged a withdrawal."),
    ("notice_type", "Rejection notice", "The lockbox rejected the packet and returned it."),
    ("notice_type", "Card produced", "The new card was produced."),
    ("notice_type", "Card mailed", "The new card was mailed."),
    ("notice_type", "Card delivered", "The carrier delivered the card."),
    ("notice_type", "Card returned undeliverable", "The card came back to USCIS undelivered."),
    ("action", "Expedite request", "A request to USCIS to expedite a pending case."),
    ("action", "Case inquiry", "A service request on a case outside the posted processing time."),
    ("action", "Biometrics reschedule request", "A request to move a biometrics appointment."),
    ("action", "Address change (AR-11)", "A change of address reported to USCIS."),
    ("action", "Withdrawal request", "A request to USCIS to withdraw a filing."),
    ("action", "Stop work", "Work stopped on the case."),
    ("action", "Notice forwarded by applicant", "A notice the applicant received and sent to Global Mobility."),
    ("action", "Card not received inquiry", "A service request on a card not received."),
]


def write_dictionary(out, stamp):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Fields"
    fill = PatternFill("solid", fgColor="DCE6F1")
    for j, h in enumerate(("File", "Field", "Definition")):
        c = ws.cell(row=1, column=j + 1, value=h)
        c.font, c.fill = Font(bold=True), fill
    for i, (f, fld, dfn) in enumerate(DICT):
        ws.cell(row=i + 2, column=1, value=f)
        ws.cell(row=i + 2, column=2, value=fld)
        ws.cell(row=i + 2, column=3, value=dfn).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 100
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("Codes")
    for j, h in enumerate(("Field", "Code", "Meaning")):
        c = ws2.cell(row=1, column=j + 1, value=h)
        c.font, c.fill = Font(bold=True), fill
    for i, (f, code, m) in enumerate(CODES):
        ws2.cell(row=i + 2, column=1, value=f)
        ws2.cell(row=i + 2, column=2, value=code)
        ws2.cell(row=i + 2, column=3, value=m)
    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 30
    ws2.column_dimensions["C"].width = 80
    ws2.freeze_panes = "A2"
    path = os.path.join(out, C.F_DICT)
    wb.save(path)
    creator, last_by, created, modified = stamp
    PP.normalize_ooxml(path, creator, last_by, created, modified, excel=True, company=C.ORG)
    return len(DICT) + len(CODES)


MANIFEST = [
    (C.F_NOTICES, "Marrowby Teale LLP case-management system, renewal matters for Sallowmere Software", "2026-12-04",
     "every notice and event logged to 4 December 2026 on the renewal matters received from 1 January 2024 or still "
     "open"),
    (C.F_REGISTER, "Global Mobility case register", "2026-12-04",
     "cases opened from 1 December 2023, and earlier cases received from 1 January 2024"),
    (C.F_ACTIONS, "Global Mobility case-action log", "2026-12-04", "actions on the register's cases to 4 December 2026"),
    (C.F_PERMITS, "HR Systems, work permit roster", "2026-12-04", "every card entered since 2021"),
    (C.F_EMPLOYEES, "HR Systems, Global Mobility program roster", "2026-12-04",
     "employees active on 4 December 2026 or terminated since 1 January 2024"),
    (C.F_DEPENDANTS, "Benefits, covered dependants", "2026-12-04", "dependants covered on 4 December 2026"),
    (C.F_TRACKER, "Global Mobility green-card tracker", "2026-12-03", "adjustment applicants in Global Mobility's care since 2024"),
    (C.F_DISPATCH, "Mailroom dispatch log", "2026-12-04", "packets sent from 2 January 2025 to 4 December 2026"),
    (C.F_POLICY, "Global Mobility", "2025-11-12", "standard GM-07, revision 4"),
    (C.F_THREAD, "Global Mobility mailbox", "2026-12-03", "the December thread on the Q1 2027 portal release"),
    (C.F_DICT, "Global Mobility", "2026-12-04", "field definitions and codes for the files in this set"),
    (C.F_POSTED, "Global Mobility", "2026-12-01", "monthly captures from January 2024 to November 2026"),
    (C.F_H1B, "Immigration team", "2026-11-20", "H-1B extensions due in 2027"),
]


def write_manifest(out, n):
    rows = []
    for f, src, d, cov in MANIFEST:
        rows.append((f, src, d, n.get(f, ""), cov))
    rows.append((C.F_MANIFEST, "Global Mobility", "2026-12-04", "", "this list"))
    return C.write_csv(os.path.join(out, C.F_MANIFEST), ("file", "source", "extract_date", "rows", "coverage"), rows)
