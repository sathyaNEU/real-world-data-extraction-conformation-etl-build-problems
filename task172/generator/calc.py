"""task172 solution, computed from the shipped files only (target/). The quote is a forecast: the 80th percentile of
the wait for the renewals received from 1 January to 31 March 2027, under the stage model the record supports.

A renewal waits for its fingerprint appointment (unless USCIS reuses fingerprints it holds), then for the officer's
review. Since renewals filed from 30 October 2025 stopped extending the card, a renewal still pending on the day the
filer's card expires has an expedite request from the team that day, and USCIS has decided it about a week later. The
record met that cut only on its slowest tail, because appointments came about four weeks after receipt; the quarter's
renewals meet appointments about sixteen weeks out, and most will reach their card's expiry undecided."""
from __future__ import annotations

import datetime as dt
import os
import sys
from collections import defaultdict

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import common as C   # noqa: E402

D_ = dt.timedelta
NEW_REGIME = (dt.date(2026, 10, 1), dt.date(2026, 11, 30))   # receipts whose appointments show the lead now set
N_SIM = 400000


def load(target):
    from openpyxl import load_workbook
    T = {}
    T["notices"] = C.read_csv(os.path.join(target, C.F_NOTICES))
    T["actions"] = C.read_csv(os.path.join(target, C.F_ACTIONS))
    T["dispatch"] = C.read_csv(os.path.join(target, C.F_DISPATCH))
    T["employees"] = C.read_csv(os.path.join(target, C.F_EMPLOYEES))
    T["dependants"] = C.read_csv(os.path.join(target, C.F_DEPENDANTS))
    T["posted"] = C.read_csv(os.path.join(target, C.F_POSTED))
    wb = load_workbook(os.path.join(target, C.F_REGISTER), read_only=True)
    old = list(wb["Cases to Mar 2025"].iter_rows(values_only=True))
    new = list(wb["Cases"].iter_rows(values_only=True))
    T["reg_old"] = [dict(zip(old[0], r)) for r in old[1:]]
    T["reg_new"] = [dict(zip(new[0], r)) for r in new[1:]]
    wb = load_workbook(os.path.join(target, C.F_PERMITS), read_only=True)
    rows = list(wb["Permits"].iter_rows(values_only=True))
    T["permits"] = [dict(zip(rows[0], r)) for r in rows[1:]]
    wb = load_workbook(os.path.join(target, C.F_TRACKER), read_only=True)
    rows = list(wb["Applicants"].iter_rows(values_only=True))
    T["tracker"] = [dict(zip(rows[0], r)) for r in rows[1:]]
    return T


def _d(v):
    """a date from an ISO string, a US text date or an Excel date"""
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    v = str(v)
    return C.parse_us(v) if "/" in v else C.parse_iso(v)


def register(T):
    """one record per case, keyed by every reference it carries (RN and GM), with the current sheet preferred"""
    cases, ref = {}, {}
    for r in T["reg_old"]:
        k = r["Case ID"]
        cases[k] = dict(id=k, type=r["Case Type"].lower(), person=r["Person ID"], ptype=r["Person Type"].lower(),
                        sponsor=r["Sponsor ID"] or "", cat=r["Category"], card_exp=_d(r["Card Expires"]),
                        opened=_d(r["Opened"]), sent=_d(r["Packet Sent"]), matter=r["Matter No."] or "",
                        receipt=r["Receipt No."] or "", received=_d(r["Received"]), status=r["Status"],
                        decided=_d(r["Decided"]), new_exp=_d(r["New Card Expires"]), closed=_d(r["Closed"]),
                        filed_by=r["Filed By"])
        ref[k] = k
    for r in T["reg_new"]:
        k = r["case_id"]
        rec = dict(id=k, type=r["case_type"], person=r["person_id"], ptype=r["person_type"], sponsor=r["sponsor_id"] or "",
                   cat=r["category"], card_exp=_d(r["card_expires"]), opened=_d(r["opened_on"]),
                   sent=_d(r["packet_sent_on"]), matter=r["matter_no"] or "", receipt=r["receipt_no"] or "",
                   received=_d(r["received_on"]), status=r["status"], decided=_d(r["decided_on"]),
                   new_exp=_d(r["new_card_expires"]), closed=_d(r["closed_on"]), filed_by=r["filed_by"],
                   legacy=r["legacy_case_id"] or "")
        cases[k] = rec
        ref[k] = k
        if rec["legacy"]:
            ref[rec["legacy"]] = k
    return cases, ref


def renewals(T, cases=None, ref=None):
    """the renewals counsel filed and USCIS receipted, from counsel's notices, with the card each renews (register)
    and the team's expedite requests (case-action log)"""
    if cases is None:
        cases, ref = register(T)
    by_receipt = defaultdict(list)
    for n in T["notices"]:
        if n["receipt_no"]:
            by_receipt[n["receipt_no"]].append(n)
    exped = defaultdict(list)
    stops = {}
    for a in T["actions"]:
        k = ref.get(a["case_id"])
        if a["action"] == "Expedite request":
            exped[k].append((_d(a["action_date"]), a["outcome"]))
        elif a["action"] == "Stop work":
            stops[k] = _d(a["action_date"])
    out = []
    for rno, ns in by_receipt.items():
        rcpt = [n for n in ns if n["notice_type"] == "Receipt notice"]
        if not rcpt:
            continue
        received = _d(rcpt[0]["received_date"])
        k = ref.get(ns[0]["client_ref"])
        reg = cases.get(k)
        appts = sorted((n for n in ns if n["notice_type"] == "Biometrics appointment notice"),
                       key=lambda n: (n["notice_date"], n["logged_on"], n["notice_id"]))
        reuse = any(n["notice_type"] == "Biometrics reuse notice" for n in ns)
        dec = [n for n in ns if n["notice_type"] in ("Approval notice", "Denial notice")]
        wd = [n for n in ns if n["notice_type"] == "Withdrawal acknowledgement"]
        decided = min(_d(n["notice_date"]) for n in dec) if dec else None
        withdrawn = bool(wd) and (decided is None or _d(wd[0]["notice_date"]) < decided)
        stop = (stops.get(k) if k else None) or (reg["closed"] if reg and reg["status"] == "Closed - no longer tracked"
                                                    else None)
        ex = sorted(exped.get(k, [])) if k else []
        out.append(dict(receipt=rno, case=k, cat=ns[0]["eligibility_category"], received=received,
                        appt=_d(appts[-1]["appointment_date"]) if appts else None,
                        appt_first=_d(appts[0]["appointment_date"]) if appts else None,
                        route="appointment" if appts else ("reuse" if reuse else None),
                        decided=None if withdrawn else decided, withdrawn=withdrawn, stop=stop,
                        card_exp=reg["card_exp"] if reg else None,
                        exp_req=ex[-1][0] if ex else None, exp_out=ex[-1][1] if ex else None,
                        approved=any(n["notice_type"] == "Approval notice" for n in dec)))
    out.sort(key=lambda r: (r["received"], r["receipt"]))
    return out


def km_points(rows):
    """Kaplan-Meier: (time, event) rows to the event times and their probability masses; mass beyond the last event
    stays at the largest time observed"""
    rows = sorted(rows)
    t = np.array([r[0] for r in rows], float)
    e = np.array([r[1] for r in rows], int)
    ut = np.unique(t[e == 1])
    S, vals, ps = 1.0, [], []
    for u in ut:
        nr = int((t >= u).sum())
        d = int(((t == u) & (e == 1)).sum())
        vals.append(u)
        ps.append(S * d / nr)
        S *= 1 - d / nr
    vals.append(float(t.max()))
    ps.append(S)
    ps = np.array(ps)
    return np.array(vals), ps / ps.sum()


def km_quantile(rows, p=0.8):
    rows = sorted(rows)
    n, S, i = len(rows), 1.0, 0
    while i < n:
        t, d, j = rows[i][0], 0, i
        while j < n and rows[j][0] == t:
            d += rows[j][1]
            j += 1
        if d:
            S *= 1 - d / (n - i)
            if S <= 1 - p + 1e-12:
                return t
        i = j
    return float("nan")


def components(rens, export=C.EXPORT, regime=NEW_REGIME, d_from=C.IFR, q_mode="km"):
    """the stage model's parts, each read from the record"""
    rec = [r for r in rens if r["received"] and r["received"] <= export and not r["withdrawn"] and r["route"]]
    pA = float(np.mean([r["route"] == "appointment" for r in rec]))
    L = np.array([(r["appt"] - r["received"]).days for r in rec if r["route"] == "appointment" and
                  regime[0] <= r["received"] <= regime[1]], float)
    qrows = []
    for r in rec:
        st = r["appt"] if r["route"] == "appointment" else r["received"]
        if st is None or st > export:
            continue
        end = min(x for x in (export, r["stop"]) if x)
        if r["exp_req"] and (r["decided"] is None or r["exp_req"] < r["decided"]):
            qrows.append(((r["exp_req"] - st).days, 0))           # the request cut the review short
        elif r["decided"] and r["decided"] <= end:
            qrows.append(((r["decided"] - st).days, 1))
        else:
            qrows.append(((end - st).days, 0))
    qrows = [q for q in qrows if q[0] >= 0]
    if q_mode == "decided":
        qrows = [q for q in qrows if q[1] == 1]
    D = np.array([(r["card_exp"] - r["received"]).days for r in rec if r["card_exp"] and r["received"] >= d_from], float)
    ex = [r for r in rens if r["exp_req"] and r["received"] >= d_from]
    with_out = [r for r in ex if r["exp_out"] in ("Accepted", "Denied")]
    accept = float(np.mean([r["exp_out"] == "Accepted" for r in with_out])) if with_out else 1.0
    E = np.array([(r["decided"] - r["exp_req"]).days for r in ex if r["exp_out"] == "Accepted" and r["decided"] and
                  r["decided"] >= r["exp_req"]], float)
    R = np.array([(r["exp_req"] - r["card_exp"]).days for r in ex if r["card_exp"]], float)
    pA_cat = {cat: float(np.mean([r["route"] == "appointment" for r in rec if r["cat"] == cat])) for cat in C.CATS}
    n_cat = {cat: sum(1 for r in rec if r["cat"] == cat) for cat in C.CATS}
    return dict(pA=pA, pA_cat=pA_cat, n_cat=n_cat, L=L, qrows=qrows, D=D, E=E, R=R, accept=accept, n_rec=len(rec))


def simulate(comp, seed=172, n=N_SIM, cap=True, use_accept=True, lead_add=None, mix=None, offsets=True):
    """the quarter's waits by simulation: the route (by category where a mix is given), the appointment lead now set,
    the review (Kaplan-Meier), and for a renewal still pending when the team's request goes in (its card's expiry plus
    the record's request offset), the decision after the record's expedite response"""
    rng = np.random.default_rng(seed)
    if mix is None:
        appt = rng.random(n) < comp["pA"]
    else:
        cats = list(mix)
        k = rng.choice(len(cats), n, p=np.array([mix[c] for c in cats], float) / sum(mix.values()))
        pa = np.array([comp["pA_cat"][c] for c in cats])[k]
        appt = rng.random(n) < pa
    L = rng.choice(comp["L"], n) if lead_add is None else lead_add(rng, n)
    vals, ps = km_points(comp["qrows"])
    Q = rng.choice(vals, n, p=ps)
    Dd = rng.choice(comp["D"], n)
    E = rng.choice(comp["E"], n)
    Rr = rng.choice(comp["R"], n) if (offsets and len(comp["R"])) else np.zeros(n)
    wn = np.where(appt, L, 0.0) + Q
    req = Dd + Rr
    if not cap:
        return wn, dict(reach=float((wn > Dd).mean()), cut=0.0)
    ok = rng.random(n) < (comp["accept"] if use_accept else 1.0)
    w = np.where((wn > req) & ok, np.minimum(wn, req + E), wn)
    return w, dict(reach=float((wn > Dd).mean()), cut=float(((wn > req) & ok & (req + E < wn)).mean()),
                   reach_appt=float((wn[appt] > Dd[appt]).mean()))


def quote(target=None, T=None):
    T = T or load(target)
    cases, ref = register(T)
    rens = renewals(T, cases, ref)
    comp = components(rens)
    mix, n_mix = forward_mix(T, cases, ref, float(np.median(comp["D"])))
    w, info = simulate(comp, mix=mix)
    q80 = float(np.quantile(w, 0.8))
    months = q80 / C.MONTH_DAYS
    return dict(q80_days=q80, months=months, quote=C.half_month(months), lead_weeks=float(np.median(comp["L"])) / 7,
                reach=info["reach"], reach_appt=info["reach_appt"], cut=info["cut"], e_days=float(np.median(comp["E"])),
                accept=comp["accept"], pA=comp["pA"], pA_cat=comp["pA_cat"], mix=mix, n_mix=n_mix, n_L=len(comp["L"]),
                n_E=len(comp["E"]), n_D=len(comp["D"]), d_median=float(np.median(comp["D"])), comp=comp, rens=rens,
                cases=cases, ref=ref, T=T)


# ------------------------------------------------------------------------------------------------------- ask A
def holders_current(T, cases):
    """employees HR's program roster shows as Active at the export, and each spouse's household employee, taken from any
    record that links the two: the roster's onboarding entry, the register's cases, the green-card tracker, or
    Benefits' coverage"""
    active = {e["employee_id"] for e in T["employees"] if e["status"] == "Active"}
    sponsor = {}
    for p in T["permits"]:
        if p.get("household_employee_id"):
            sponsor.setdefault(p["holder_id"], p["household_employee_id"])
    for c in cases.values():
        if c["ptype"] == "spouse" and c["sponsor"]:
            sponsor.setdefault(c["person"], c["sponsor"])
    for t in T["tracker"]:
        if t["relationship"] == "Spouse":
            sponsor.setdefault(t["applicant_id"], t["principal_employee_id"])
    enroll = {e["benefits_enrollment_id"]: e["employee_id"] for e in T["employees"]}
    for d in T["dependants"]:
        if d["relationship"] in ("Spouse", "Domestic partner") and d["benefits_enrollment_id"] in enroll:
            sponsor.setdefault(d["dependant_id"], enroll[d["benefits_enrollment_id"]])
    return active, sponsor


def cards_held(T, cases, ref, sources=("roster", "register", "counsel", "tracker")):
    """every work permit card on record, one per holder and approval date: the roster's entries, the register's approved
    filings (renewals and initial filings), counsel's approval notices and the tracker's (c)(9) cards; the category is
    the one USCIS approved, from a filing record wherever one holds the card, the roster's entry only where none does"""
    cards = {}

    def put(holder, vf, vt, cat, src, filing):
        if not (holder and vf and vt):
            return
        key = (holder, vf)
        cur = cards.get(key)
        if cur is None:
            cards[key] = dict(holder=holder, vf=vf, vt=vt, cat=cat, src=src, filing=filing, srcs={src})
            return
        cur["srcs"].add(src)
        if filing and not cur["filing"]:
            cur.update(cat=cat, filing=True)
    if "roster" in sources:
        for p in T["permits"]:
            put(p["holder_id"], _d(p["valid_from"]), _d(p["valid_to"]), p["category"], "roster", False)
    if "register" in sources:
        for c in cases.values():
            if c["status"] == "Approved" and c["decided"] and c["new_exp"]:
                put(c["person"], c["decided"], c["new_exp"], c["cat"], "register", True)
    if "counsel" in sources:
        for n in T["notices"]:
            if n["notice_type"] == "Approval notice":
                c = cases.get(ref.get(n["client_ref"]))
                if c is not None:
                    put(c["person"], _d(n["valid_from"]), _d(n["valid_to"]), n["eligibility_category"], "counsel", True)
    if "tracker" in sources:
        for t in T["tracker"]:
            if t["ead_approved"] and t["ead_valid_to"]:
                put(t["applicant_id"], _d(t["ead_approved"]), _d(t["ead_valid_to"]), "(c)(9)", "tracker", True)
    return [c for c in cards.values() if c["vf"] <= C.EXPORT]


def latest_cards(T, cases, ref, sources=("roster", "register", "counsel", "tracker"), links="all", merge=True):
    """GM-07 1.3 and 1.4: each current person's work permit, the card USCIS approved for the person most recently; a
    person is current while HR's program roster shows the employee (or, for a spouse, the employee whose household she
    belongs to) as Active. One person may carry two IDs: a spouse keeps her Global Mobility person ID (MP-) on the
    team's records and is known to Benefits, and so to the benefits portal, by a dependant ID once she joins the plan;
    an employee has one spouse, so a spouse is the household she belongs to (merge=False keys persons by ID instead)"""
    active, sponsor = holders_current(T, cases)
    if links == "benefits":
        enroll = {e["benefits_enrollment_id"]: e["employee_id"] for e in T["employees"]}
        sponsor = {d["dependant_id"]: enroll[d["benefits_enrollment_id"]] for d in T["dependants"]
                   if d["relationship"] != "Child" and d["benefits_enrollment_id"] in enroll}
    emp_ids = {e["employee_id"] for e in T["employees"]}
    best = {}
    for c in cards_held(T, cases, ref, sources):
        h = c["holder"]
        if h in emp_ids:
            person, cur = h, h in active
        else:
            e = sponsor.get(h)
            person, cur = (("spouse of", e) if merge else h), e in active
        if not cur:
            continue
        b = best.get(person)
        if b is None or c["vf"] > b["vf"]:
            best[person] = dict(c)
        elif c["vf"] == b["vf"] and c["filing"] and not b["filing"]:
            b.update(cat=c["cat"], filing=True)      # one card under two IDs: the approval's category stands
    return best


def ask_a(T, cases, ref, **kw):
    out = {(y, m, cat): 0 for (y, m) in C.ASK_A_MONTHS for cat in C.CATS}
    for c in latest_cards(T, cases, ref, **kw).values():
        key = (c["vt"].year, c["vt"].month, c["cat"])
        if key in out:
            out[key] += 1
    return out


def forward_mix(T, cases, ref, d_median):
    """the quarter's renewals by category: the current work permits whose renewal reaches USCIS in the quarter (the
    card's last day less the record's median card gap falls from 1 January to 31 March 2027), leaving out the (c)(9)
    cards of applicants whose adjustment the tracker shows approved"""
    lpr = set()
    for t in T["tracker"]:
        if t["i485_status"] == "Approved":
            lpr.add(t["applicant_id"] if t["relationship"] == "Principal" else ("spouse of", t["principal_employee_id"]))
    n = {cat: 0 for cat in C.CATS}
    for person, c in latest_cards(T, cases, ref).items():
        if c["cat"] == "(c)(9)" and person in lpr:
            continue
        if C.Q_START <= c["vt"] - D_(days=int(round(d_median))) <= C.Q_END:
            n[c["cat"]] += 1
    tot = sum(n.values())
    return {cat: n[cat] / tot for cat in C.CATS}, n


# ------------------------------------------------------------------------------------------------------- ask B
def rejections(T, cases, ref):
    """the lockbox's rejections of renewal packets: counsel's rejection notices, and the rejection notices applicants
    forwarded that counsel never had; one per case and notice date"""
    rej = {}
    for n in T["notices"]:
        if n["notice_type"] == "Rejection notice":
            k = ref.get(n["client_ref"])
            rej[(k, _d(n["notice_date"]))] = "counsel"
    for a in T["actions"]:
        if a["action"] == "Notice forwarded by applicant" and a["detail"] == "Rejection notice":
            k = ref.get(a["case_id"])
            rej.setdefault((k, _d(a["notice_date"])), "applicant")
    return rej


def ask_b(T, cases, ref):
    rej = rejections(T, cases, ref)
    receipts = {}
    for n in T["notices"]:
        if n["notice_type"] == "Receipt notice":
            k = ref.get(n["client_ref"])
            d = _d(n["received_date"])
            receipts.setdefault(k, []).append(d)
    for c in cases.values():
        if c["type"] == "renewal" and c["received"]:
            receipts.setdefault(c["id"], []).append(c["received"])
    out = {}
    for label, a, b in C.ASK_B_QUARTERS:
        n_rej, n_ref = 0, 0
        for (k, nd), _src in rej.items():
            if not (a <= nd <= b):
                continue
            n_rej += 1
            later = sorted(x for x in receipts.get(k, []) if x and x > nd)
            if later and (later[0] - nd).days <= C.REFILE_DAYS:
                n_ref += 1
        out[label] = (n_rej, n_ref)
    return out


def solve(target):
    Q = quote(target)
    T, cases, ref = Q["T"], Q["cases"], Q["ref"]
    A = ask_a(T, cases, ref)
    B = ask_b(T, cases, ref)
    return Q, A, B


if __name__ == "__main__":
    Q, A, B = solve(sys.argv[1])
    print("quote %.1f months (q80 %.1f days = %.3f months); lead %.2f weeks (n %d); E median %.1f days (n %d); "
          "accept %.3f; reach %.3f (appointment route %.3f); cut %.3f; pA %.3f by category %s; mix %s; D median %.0f (n %d)"
          % (Q["quote"], Q["q80_days"], Q["months"], Q["lead_weeks"], Q["n_L"], Q["e_days"], Q["n_E"], Q["accept"],
             Q["reach"], Q["reach_appt"], Q["cut"], Q["pA"], {k: round(v, 3) for k, v in Q["pA_cat"].items()},
             Q["n_mix"], Q["d_median"], Q["n_D"]))
    for (y, m), in [((y, m),) for (y, m) in C.ASK_A_MONTHS]:
        print("  %d-%02d  (c)(9) %4d  (c)(26) %4d" % (y, m, A[(y, m, "(c)(9)")], A[(y, m, "(c)(26)")]))
    for k, v in B.items():
        print("  %s  rejected %3d  refiled within 30 days %3d" % (k, v[0], v[1]))
