"""The pack's record files as each system exports them: the location vendor's timing-point departures (first release
and revised), the ticketing system's journeys, the timetable feeds, the corridor workbook, the stop changes log, the
operator's pilot timesheets and staff list, and the agency's ledger. Each writer returns its row count."""
import csv
import datetime as dt
import hashlib
import io
import zipfile

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

import common as K
import world as W

RELOCATED = None   # set by relocation(net): (old stop number, new stop number)


def _csv(path, header, rows, delim=","):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=delim, lineterminator="\r\n")
        w.writerow(header)
        w.writerows(rows)
    return len(rows)


def _sheet(ws, header, rows, widths, date_cols=(), freeze=True, fmt=None):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE7EE")
        c.alignment = Alignment(vertical="top", wrap_text=True)
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    for c in date_cols:
        for cell in ws[openpyxl.utils.get_column_letter(c + 1)][1:]:
            if isinstance(cell.value, (dt.date, dt.datetime)):
                cell.number_format = "dd/mm/yyyy"
    for c, f in (fmt or {}).items():
        for cell in ws[openpyxl.utils.get_column_letter(c + 1)][1:]:
            cell.number_format = f
    if freeze:
        ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def hms(m):
    s = int(round(m * 60))
    return "%02d:%02d:%02d" % (s // 3600, (s // 60) % 60, s % 60)


def stamp(day, m):
    s = int(round(m * 60))
    t = dt.datetime(day.year, day.month, day.day) + dt.timedelta(seconds=s)
    return t.strftime("%Y-%m-%dT%H:%M:%S") + K.TZ


def in_extract(m):
    return any(a <= m < b for a, b in K.EXTRACT)


def relocation(net):
    """Route 22's outbound stop after Brewery Lane moved 80 metres for roadworks from Monday 17 August and took a new
    stop number; the old number was retired."""
    global RELOCATED
    cor = net["corridor"][("22", "out")]
    old = cor[5]
    RELOCATED = (old, old + 88)
    return RELOCATED


def stop_number(code, day):
    if RELOCATED and code == RELOCATED[0] and day >= K.RELOCATION:
        return RELOCATED[1]
    return code


def trip_meta(trip):
    head = trip.split("-")[0]
    rt = "".join(ch for ch in head if ch.isdigit())
    d = "out" if head[len(rt)] == "O" else "in"
    return rt, d, head[len(rt) + 1:]


def headsign(net, rt, d, v):
    if d == "in":
        return "City"
    if v:
        return W.NET[rt]["branches"][v]["name"]
    return W.NET[rt]["tps"][-1]


# ------------------------------------------------------------------------------------------------- departures
def departures(net, S, seed=W.SEED):
    """The vendor's revised departures at timing points, one row a trip a timing point (not the trip's last stop),
    in the extract windows; a bus that reopened its doors at a timing point left two records, the later one the
    time it left. Returns (revised rows, first-release rows)."""
    rng = np.random.default_rng([seed, 31])
    fleet = {}
    rows = []
    for (day, trip), arr in sorted(S["deps"].items(), key=lambda kv: (kv[0][0], kv[0][1])):
        rt, d, v = trip_meta(trip)
        pat = net["patterns"][(rt, d, v)]
        sch = S["sched"][(day, trip)]
        veh = fleet.setdefault((day, trip), 1100 + int(rng.integers(0, 420)))
        for i, s in enumerate(pat[:-1]):
            if not net["stops"][s]["tp"] or np.isnan(arr[i]) or not in_extract(arr[i]):
                continue
            rows.append([day, rt, "Outbound" if d == "out" else "Inbound", trip, headsign(net, rt, d, v), veh,
                         stop_number(s, day), net["stops"][s]["name"], hms(sch[i]), float(arr[i]), 0])
    # doubled records where the doors reopened: an earlier record 15 to 70 seconds before the bus left
    extra = []
    for r in rows:
        if rng.random() < 0.006:
            e = list(r)
            e[9] = r[9] - rng.uniform(15, 70) / 60
            e[10] = 1
            extra.append(e)
    allr = rows + extra
    allr.sort(key=lambda r: (r[0], r[9], r[3]))
    first = [r for r in allr if not (r[10] == 0 and rng.random() < 0.011)]
    return allr, first


def write_departures(path, rows, release):
    out = []
    for n, r in enumerate(rows, 1):
        day = r[0]
        rid = "%s-%s-%06d" % (release, day.strftime("%m%d"), n)
        out.append([day.isoformat(), r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8], stamp(day, r[9]), rid])
    return _csv(path, ["service_date", "route", "direction", "trip_id", "destination", "vehicle", "timing_point",
                       "timing_point_name", "scheduled_time", "departure_time", "record_id"], out)


# ------------------------------------------------------------------------------------------------- journeys
def card_id(n):
    return hashlib.sha1(("pellcard-%d" % n).encode()).hexdigest()[:14].upper()


def write_journeys(path, J):
    rows = []
    sel = [j for j in J if in_extract(j["on"])]
    sel.sort(key=lambda j: (j["day"], j["on"], j["trip"]))
    for n, j in enumerate(sel, 1):
        day = j["day"]
        off_stop = stop_number(j["alight"], day) if j["alight"] is not None else ""
        rows.append(["J%s%07d" % (day.strftime("%m%d"), n), card_id(j["card"]), j["fare"], j["route"], j["trip"],
                     stop_number(j["board"], day), stamp(day, j["on"]), off_stop,
                     stamp(day, j["off"]) if j["off"] is not None else "",
                     "Complete" if j["alight"] is not None else "Default fare"])
    return _csv(path, ["journey_id", "card_id", "fare_type", "route", "trip_id", "tag_on_stop", "tag_on_time",
                       "tag_off_stop", "tag_off_time", "fare_status"], rows)


# ------------------------------------------------------------------------------------------------- timetable feeds
def _zip(path, files):
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, text in files:
            info = zipfile.ZipInfo(name, date_time=(2026, 7, 17, 9, 30, 0) if "07-20" in path else (2026, 8, 13, 11, 5, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            z.writestr(info, text)


def _txt(header, rows):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(header)
    w.writerows(rows)
    return buf.getvalue()


def write_gtfs(path, net, tt, feed):
    a, b = K.FEED1 if feed == 1 else K.FEED2
    sid = "WD-%s" % a.strftime("%m%d")
    day = a
    stops_used = set()
    trips, times = [], []
    for x in tt:
        rt, d, v = x["route"], x["direction"], x["variant"]
        pat = net["patterns"][(rt, d, v)]
        trips.append([rt, sid, x["trip"], headsign(net, rt, d, v), 0 if d == "out" else 1,
                      "%s-%s%s" % (rt, d.upper(), v)])
        for i, s in enumerate(pat):
            code = stop_number(s, day)
            stops_used.add(code)
            t = hms(x["times"][i])
            times.append([x["trip"], t, t, code, i + 1, 1 if net["stops"][s]["tp"] else 0])
    stops = []
    for code in sorted(stops_used):
        src = RELOCATED[0] if RELOCATED and code == RELOCATED[1] else code
        st = net["stops"][src]
        lat, lon = _coords(src)
        if RELOCATED and code == RELOCATED[1]:
            lat += 0.0007
        stops.append([code, code, st["name"], "%.6f" % lat, "%.6f" % lon])
    routes = [[rt, "PM", rt, "%s" % W.NET[rt]["road"], 3, "1f5a8c" if rt in K.BRANCHED else "2e7d4f"] for rt in K.ROUTES]
    files = [
        ("agency.txt", _txt(["agency_id", "agency_name", "agency_url", "agency_timezone", "agency_lang"],
                            [["PM", K.AGENCY, "https://www.pellbrookmetro.example", "Australia/Brisbane", "en"]])),
        ("feed_info.txt", _txt(["feed_publisher_name", "feed_publisher_url", "feed_lang", "feed_start_date",
                                "feed_end_date", "feed_version"],
                               [[K.AGENCY, "https://www.pellbrookmetro.example", "en", a.strftime("%Y%m%d"),
                                 b.strftime("%Y%m%d"), "%s.1" % a.strftime("%Y%m%d")]])),
        ("calendar.txt", _txt(["service_id", "monday", "tuesday", "wednesday", "thursday", "friday", "saturday",
                               "sunday", "start_date", "end_date"],
                              [[sid, 1, 1, 1, 1, 1, 0, 0, a.strftime("%Y%m%d"), b.strftime("%Y%m%d")]])),
        ("routes.txt", _txt(["route_id", "agency_id", "route_short_name", "route_long_name", "route_type",
                             "route_color"], routes)),
        ("stops.txt", _txt(["stop_id", "stop_code", "stop_name", "stop_lat", "stop_lon"], stops)),
        ("trips.txt", _txt(["route_id", "service_id", "trip_id", "trip_headsign", "direction_id", "shape_id"], trips)),
        ("stop_times.txt", _txt(["trip_id", "arrival_time", "departure_time", "stop_id", "stop_sequence",
                                 "timepoint"], times)),
    ]
    _zip(path, files)
    return len(times)


def _coords(code):
    r = np.random.default_rng([W.SEED, 41, int(code)])
    return -27.47 + r.uniform(-0.18, 0.18), 153.02 + r.uniform(-0.2, 0.2)


# ------------------------------------------------------------------------------------------------- corridors
def write_corridors(path, net):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Corridor stops"
    rows = []
    for rt in K.ROUTES:
        for d in ("out", "in"):
            for i, s in enumerate(net["corridor"][(rt, d)], 1):
                rows.append([rt, W.NET[rt]["road"], "Outbound" if d == "out" else "Inbound", i, s,
                             net["stops"][s]["name"], "Yes" if net["stops"][s]["tp"] else "No"])
    _sheet(ws, ["Route", "Corridor", "Direction", "Sequence", "Stop number", "Stop name", "Timing point"], rows,
           [7, 16, 11, 9, 12, 40, 12])
    ws2 = wb.create_sheet("About")
    notes = [
        ["Frequent corridors, July 2026"],
        ["Network Planning, Pellbrook Metro. Issued 6 July 2026 for the July timetable."],
        [None],
        ["A frequent corridor is the run of stops on a frequent route where a bus comes at least every ten minutes "
         "from 7 am to 7 pm on weekdays. On routes 7, 14, 31 and 52 it runs from the City Interchange to the junction "
         "where the route's two branches part; on routes 3, 9, 22 and 40 it is the whole route."],
        ["Stops are listed in travel order for each direction. Timing points are the stops the timetable times "
         "buses at."],
    ]
    for r in notes:
        ws2.append(r)
    ws2.column_dimensions["A"].width = 110
    for row in ws2.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws2["A1"].font = Font(bold=True, size=12)
    wb.save(path)
    return len(rows)


def write_stop_changes(path, net):
    old, new = RELOCATED
    nm = net["stops"][old]["name"]
    def nm_(code):
        return net["stops"][code]["name"]
    rows = [
        ["SC-26-014", "2026-03-09", 209109, "", nm_(209109), "Renamed", "9", "Stop name corrected on signs and feed"],
        ["SC-26-021", "2026-05-18", 252110, "", nm_(252110), "Shelter replaced", "52", "No change to stop number"],
        ["SC-26-033", "2026-08-17", old, new, nm, "Relocated", "22",
         "Moved 80 m north for the Mardale Road roadworks; new stop number from the first service"],
        ["SC-26-034", "2026-08-17", 240122, "", nm_(240122), "Temporary closure", "40",
         "Closed 9 pm to 5 am for night works; weekday services unaffected"],
        ["SC-26-041", "2026-09-28", 231114, "", nm_(231114), "Accessibility upgrade", "31", "Raised platform"],
    ]
    return _csv(path, ["change_id", "effective_date", "stop_number", "new_stop_number", "location", "change",
                       "routes", "note"], rows)


# ------------------------------------------------------------------------------------------------- the pilot
def write_staff(path, P):
    rows = [[s["emp"], s["app"], s["surname"], s["given"], s["role"], s["depot"], s["start"].isoformat()]
            for s in sorted(P["staff"], key=lambda s: s["emp"])]
    return _csv(path, ["employee_number", "roster_app_user", "surname", "given_name", "position", "depot",
                       "start_date"], rows)


def timesheet_rows(P, seed=W.SEED):
    """The operator's export of supervisor shifts, February to June 2026, stacking its payroll timesheets (to Friday
    15 May) and its rostering app (from Monday 4 May); in the two cut-over weeks every shift went into both."""
    rng = np.random.default_rng([seed, 51])
    rows = []
    a, b = K.CUTOVER
    for s in P["shifts"]:
        day = s["day"]
        (h0, m0), (h1, m1) = W.SHIFTS[s["peak"]]
        a0, a1 = h0 * 60 + m0, h1 * 60 + m1
        part = s.get("part", 0)
        if part == 1:
            a1 = a0 + 75
        if part == 2:
            a0 = a0 + 75
        hours = (a1 - a0) / 60
        if s.get("route"):
            post, loc = W.POSTS[s["route"]]
            app_route = s["route"]
        else:
            post, loc = s["depot"], W.DEPOT_POSTS[s["depot"]]
            app_route = ""
        payroll = day <= b
        app = day >= a
        if payroll and s["status"] == "Worked":
            rows.append(dict(src="Payroll", day=day, emp=s["who"]["emp"], post=post, date_text=day.strftime("%d/%m/%Y"),
                             shift=s["peak"], hours="%.2f" % hours))
        if app:
            st = dt.datetime(day.year, day.month, day.day) + dt.timedelta(minutes=a0)
            en = dt.datetime(day.year, day.month, day.day) + dt.timedelta(minutes=a1)
            who = s["who"]["app"] if s["who"] else ""
            rows.append(dict(src="Roster app", day=day, user=who, route=app_route, location=loc,
                             start=st.strftime("%Y-%m-%dT%H:%M") + K.TZ, end=en.strftime("%Y-%m-%dT%H:%M") + K.TZ,
                             status=s["status"]))
    order = sorted(range(len(rows)), key=lambda i: (rows[i]["day"], rows[i]["src"], rng.random()))
    return [rows[i] for i in order]


def write_timesheets(path, rows):
    out = []
    for r in rows:
        if r["src"] == "Payroll":
            out.append(["Payroll", r["emp"], r["post"], r["date_text"], r["shift"], r["hours"], "", "", "", "", "", ""])
        else:
            out.append(["Roster app", "", "", "", "", "", r["user"], r["route"], r["location"], r["start"], r["end"],
                        r["status"]])
    return _csv(path, ["source_system", "employee_number", "post_code", "work_date", "shift", "hours",
                       "app_user", "route", "location", "shift_start", "shift_end", "status"], out)


def ledger_rows(P, T, seed=W.SEED):
    """The agency's general ledger lines on the network performance cost centres, July 2025 to September 2026: the
    operator's monthly pilot invoices by route (posted the month after), the coordination fee, the June charges
    accrued at the 30 June year end (each route's cost centre) and reversed in July in one line at the central cost
    centre, June's invoices and fee posted in July, and the centre's other running costs."""
    rng = np.random.default_rng([seed, 61])
    rows = []
    inv = iter(range(5107, 9999))
    names = {1: "Jan", 2: "Feb", 3: "Mar", 4: "Apr", 5: "May", 6: "Jun", 7: "Jul", 8: "Aug", 9: "Sep", 10: "Oct",
             11: "Nov", 12: "Dec"}

    def post(day, cc, acct, amt, doc, narr):
        rows.append([day, cc, acct, round(amt, 2), doc, narr])

    for m in (2, 3, 4, 5, 6):
        when = dt.date(2026, m + 1, int(rng.integers(9, 17))) if m < 6 else dt.date(2026, 7, int(rng.integers(13, 18)))
        no = "PBC-%d" % next(inv)
        for rt in K.ROUTES:
            n = T["post_days"][m].get(rt, 0)
            if n:
                post(when, "NP%s" % rt.zfill(2), "6421", n * K.POST_DAY, no,
                     "PBC headway supervision %s-26 route %s" % (names[m], rt))
        fee_day = when if m < 6 else dt.date(2026, 7, int(rng.integers(20, 25)))
        post(fee_day, "NP00", "6422", K.FEE[m], "PBC-%d" % next(inv), "PBC headway pilot coordination %s-26" % names[m])
    acc = 0.0
    for rt in K.ROUTES:
        n = T["post_days"][6].get(rt, 0)
        if n:
            post(dt.date(2026, 6, 30), "NP%s" % rt.zfill(2), "6421", n * K.POST_DAY, "JNL-260630-11",
                 "PBC headway supervision Jun-26 route %s" % rt)
            acc += n * K.POST_DAY
    post(dt.date(2026, 7, 1), "NP00", "6421", -acc, "JNL-260701-02", "Reverse 30 June journal JNL-260630-11")
    # the central cost centre's other running costs
    for y, m in [(2025, x) for x in range(7, 13)] + [(2026, x) for x in range(1, 10)]:
        d0 = dt.date(y, m, int(rng.integers(3, 9)))
        post(d0, "NP00", "6310", 4180.00, "VND-%d" % int(rng.integers(20000, 29999)), "Location data licence %s-%d" %
             (names[m], y % 100))
        if m in (3, 8) or (y == 2025 and m == 11):
            post(dt.date(y, m, int(rng.integers(18, 27))), "NP00", "6450", float(np.round(rng.uniform(7800, 12400), 2)),
                 "VND-%d" % int(rng.integers(20000, 29999)), "On-board survey fieldwork %s-%d" % (names[m], y % 100))
    rows.sort(key=lambda r: (r[0], r[1], r[4]))
    return rows


CC_NAMES = {"NP00": "Network Performance (central)"}
CC_NAMES.update({"NP%s" % rt.zfill(2): "Network Performance, route %s" % rt for rt in K.ROUTES})
ACCOUNTS = {"6421": "Contracted supervision", "6422": "Contract management fees", "6310": "Data licences",
            "6450": "Surveys and research"}


def fy_period(day):
    fy = day.year + 1 if day.month >= 7 else day.year
    p = day.month - 6 if day.month >= 7 else day.month + 6
    return "FY%d-P%02d" % (fy % 100, p)


def write_ledger(path, rows):
    out = [[r[0].isoformat(), fy_period(r[0]), r[1], CC_NAMES[r[1]], r[2], ACCOUNTS[r[2]], "%.2f" % r[3], r[4], r[5]]
           for r in rows]
    return _csv(path, ["posting_date", "period", "cost_centre", "cost_centre_name", "account", "account_name",
                       "amount", "document", "narration"], out)


# ------------------------------------------------------------------------------------------------- distractors
FEEDBACK_CATS = [("Bus did not come or long wait", 0.24), ("Buses arrived together", 0.11), ("Running late", 0.16),
                 ("Driver", 0.14), ("Crowding", 0.12), ("Cleanliness", 0.05), ("Tickets and fares", 0.10),
                 ("Stop or shelter", 0.08)]


def write_feedback(path, seed=W.SEED):
    """The customer contact centre's feedback cases on the frequent routes, Term 3 2026 (20 July to 25 September)."""
    rng = np.random.default_rng([seed, 71])
    weight = {"3": 1.0, "7": 1.1, "9": 0.7, "14": 1.3, "22": 0.8, "31": 0.9, "40": 1.4, "52": 0.9}
    rts = list(weight)
    p = np.array([weight[r] for r in rts]) / sum(weight.values())
    cats = [c for c, _ in FEEDBACK_CATS]
    cp = np.array([w for _, w in FEEDBACK_CATS])
    rows = []
    day0 = dt.date(2026, 7, 20)
    n = 318
    for i in range(n):
        d = day0 + dt.timedelta(days=int(rng.integers(0, 68)))
        rt = rts[int(rng.choice(len(rts), p=p))]
        c = cats[int(rng.choice(len(cats), p=cp / cp.sum()))]
        ch = ["Phone", "Web form", "Email", "App"][int(rng.choice(4, p=[0.31, 0.37, 0.12, 0.20]))]
        st = "Closed" if d < dt.date(2026, 9, 14) or rng.random() < 0.6 else "Open"
        rows.append(["CF-%06d" % (412000 + i * 7 + int(rng.integers(0, 6))), d, rt, c, ch, st])
    rows.sort(key=lambda r: (r[1], r[0]))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cases"
    _sheet(ws, ["Case", "Received", "Route", "Category", "Channel", "Status"], rows, [12, 12, 8, 30, 11, 9],
           date_cols=(1,))
    wb.save(path)
    return len(rows)
