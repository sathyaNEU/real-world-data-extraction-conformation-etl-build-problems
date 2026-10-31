"""The pack's documents: the Foundation's grant terms, the hospital's hand hygiene program summary and its hospital-onset
bloodstream infection surveillance protocol (PDF); the ward and area reference workbook (XLSX)."""
from __future__ import annotations

import csv

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                    # noqa: E402
from reportlab.lib.pagesizes import A4                                              # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                     # noqa: E402
from reportlab.lib.units import cm                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                            # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                        # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle   # noqa: E402

import common as K   # noqa: E402
import records as R  # noqa: E402
import side as SD    # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("Sans", "LiberationSans-Regular"), ("Sans-B", "LiberationSans-Bold"), ("Sans-I", "LiberationSans-Italic"),
              ("Serif", "LiberationSerif-Regular"), ("Serif-B", "LiberationSerif-Bold")):
    pdfmetrics.registerFont(TTFont(nm, FONT_DIR + f + ".ttf"))
pdfmetrics.registerFontFamily("Sans", normal="Sans", bold="Sans-B", italic="Sans-I", boldItalic="Sans-B")
pdfmetrics.registerFontFamily("Serif", normal="Serif", bold="Serif-B", italic="Serif", boldItalic="Serif-B")
INK, GREY, RULE = colors.HexColor("#1b2430"), colors.HexColor("#56616c"), colors.HexColor("#b8c0c8")
PLUM = colors.HexColor("#5a2a52")
TEAL = colors.HexColor("#1d5a57")
NAVY = colors.HexColor("#173b5c")


def styles(face="Sans", accent=NAVY):
    b = face + "-B"
    body = ParagraphStyle("b", fontName=face, fontSize=9.8, leading=13.4, textColor=INK, spaceAfter=5)
    return dict(
        body=body,
        h1=ParagraphStyle("h1", fontName=b, fontSize=14.5, leading=18, textColor=accent, spaceAfter=2),
        h2=ParagraphStyle("h2", fontName=b, fontSize=10.6, leading=14, textColor=INK, spaceBefore=8, spaceAfter=2,
                          keepWithNext=1),
        meta=ParagraphStyle("m", fontName=face, fontSize=8.4, leading=11.2, textColor=GREY, spaceAfter=9),
        cell=ParagraphStyle("c", fontName="Sans", fontSize=7.9, leading=9.9, textColor=INK),
        cellb=ParagraphStyle("cb", fontName="Sans-B", fontSize=7.9, leading=9.9, textColor=INK),
        clause=ParagraphStyle("p", parent=body, leftIndent=28, firstLineIndent=-28),
        bullet=ParagraphStyle("u", parent=body, leftIndent=14, firstLineIndent=-9, spaceAfter=3))


def _frame(left, right):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(GREY)
        canvas.setFont("Sans", 7.4)
        canvas.drawString(2.0 * cm, A4[1] - 1.25 * cm, left)
        canvas.drawRightString(A4[0] - 2.0 * cm, A4[1] - 1.25 * cm, right)
        canvas.drawRightString(A4[0] - 2.0 * cm, 1.1 * cm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _table(S, rows, widths, right=()):
    data = [[Paragraph(str(c), S["cellb"] if i == 0 else S["cell"]) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=[w * cm for w in widths], repeatRows=1, hAlign="LEFT")
    st = [("LINEBELOW", (0, 0), (-1, 0), .6, INK), ("LINEBELOW", (0, -1), (-1, -1), .4, RULE),
          ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 1.6),
          ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6)]
    for c in right:
        st.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
    t.setStyle(TableStyle(st))
    return t


def _clause(S, num, text):
    return Paragraph("<b>%s</b>&nbsp;&nbsp;%s" % (num, text), S["clause"])


def _doc(path, title, author, subject, left, right):
    return SimpleDocTemplate(path, pagesize=A4, leftMargin=2.0 * cm, rightMargin=2.0 * cm, topMargin=2.0 * cm,
                             bottomMargin=1.9 * cm, title=title, author=author, subject=subject, creator=author,
                             producer=author), _frame(left, right)


# ------------------------------------------------------------------------------------------------ the grant's terms
def write_terms(path):
    S = styles("Serif", PLUM)
    doc, fr = _doc(path, "%s: terms" % K.GRANT, K.FOUNDATION, "Grant terms", K.FOUNDATION, "%s: terms" % K.GRANT)
    st = [Paragraph(K.GRANT, S["h1"]),
          Paragraph("Terms adopted by the board of the %s on 22 September 2026. Grants manager: %s." %
                    (K.FOUNDATION, K.PEOPLE["grants"]), S["meta"]),
          Paragraph("1. Purpose", S["h2"]),
          _clause(S, "1.1", "The Foundation funds a hand hygiene coaching nurse, one full-time position for twelve "
                            "months, on each ward of %s where staff performed hand hygiene least reliably. The "
                            "coach works on the ward's day shifts beside its nurses, doctors and other staff, "
                            "coaching hand hygiene at the bedside." % K.HOSPITAL),
          _clause(S, "1.2", "Coaching nurses start on 1 February 2027. The positions are additional to the ward's "
                            "nursing establishment and are not counted in its nursing hours."),
          Paragraph("2. The wards funded", S["h2"]),
          _clause(S, "2.1", "A ward is funded when its staff performed hand hygiene at fewer than 80 per cent of the "
                            "hand hygiene moments that arose on its AM and PM shifts in the 2025-26 audit year, 1 July "
                            "2025 to 30 June 2026. The 80 per cent line is the national benchmark for hand hygiene "
                            "compliance."),
          _clause(S, "2.2", "The moments are the five moments for hand hygiene, as the national program defines "
                            "them, on every AM and PM shift of the year. A ward's moments on a shift are taken in "
                            "proportion to the care hours its patients required on that shift, as the hospital's "
                            "patient acuity system assessed them."),
          _clause(S, "2.3", "Hand hygiene was performed at a moment when it was done by hand rub or hand wash. The "
                            "hospital's gold-auditor sessions are the record of the moments observed."),
          _clause(S, "2.4", "The Foundation funds at most six coaching nurses in 2027, one on each funded ward."),
          Paragraph("3. Evaluation", S["h2"]),
          _clause(S, "3.1", "The Foundation evaluates the grant on three outcome measures for every ward, funded or "
                            "not: hospital-onset bloodstream infections per 10,000 occupied bed days, as the "
                            "hospital's infection prevention and control service defines them, and alcohol-based hand "
                            "rub and liquid hand soap used per 1,000 occupied bed days, as the hospital's hand hygiene "
                            "program measures them."),
          _clause(S, "3.2", "The application gives the three measures for every ward for the 2025-26 audit year, the "
                            "grant's baseline."),
          Paragraph("4. Timetable", S["h2"]),
          _clause(S, "4.1", "The Foundation's board approves its 2027 grants budget on 8 December 2026."),
          _clause(S, "4.2", "Funded wards report their coaching activity to the grants manager each quarter.")]
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ------------------------------------------------------------------------------------------------ the program
def write_program(path):
    S = styles("Sans", TEAL)
    doc, fr = _doc(path, "Hand Hygiene Program 2025-26", K.PEOPLE["ipc"], "Infection prevention and control",
                   "%s, Infection Prevention and Control" % K.HOSPITAL, "Hand Hygiene Program 2025-26")
    moments = ["Moment 1, before touching a patient", "Moment 2, before a procedure",
               "Moment 3, after a procedure or body fluid exposure risk", "Moment 4, after touching a patient",
               "Moment 5, after touching a patient's surroundings"]
    st = [Paragraph("Hand Hygiene Program 2025-26", S["h1"]),
          Paragraph("%s, Infection Prevention and Control Service. Program summary for the 2025-26 audit year, "
                    "issued 14 July 2026 by %s, Clinical Nurse Consultant." % (K.HOSPITAL, K.PEOPLE["ipc"]), S["meta"]),
          Paragraph("1. The program", S["h2"]),
          Paragraph("The hospital takes part in the national hand hygiene program. Compliance is measured by direct "
                    "observation against the five moments for hand hygiene, by gold auditors trained and validated "
                    "under the national program, and reported against the national benchmark of 80 per cent. The "
                    "audit year runs from 1 July to 30 June.", S["body"]),
          Paragraph("2. The five moments", S["h2"])]
    st += [Paragraph("&bull;&nbsp;&nbsp;" + m, S["bullet"]) for m in moments]
    st += [Paragraph("3. Gold auditors and sessions", S["h2"]),
           Paragraph("Each ward and department has its own gold auditors, nurses and other clinicians from its staff "
                     "who have completed gold auditor training and annual validation. In a session the auditor "
                     "observes the moments that arise in view for about twenty minutes and records each one: which "
                     "moment it was, the health care worker group of the person it belonged to, and whether hand "
                     "hygiene was performed by hand rub or by hand wash or was missed. Each ward aims to audit at "
                     "least 100 moments a month.", S["body"]),
           Paragraph("Sessions are entered in the hand hygiene audit application on a tablet at the end of the "
                     "session. The application was unavailable from 9 to 20 March 2026; sessions in those days were "
                     "recorded on the paper form and entered by their auditors when it returned.", S["body"]),
           Paragraph("4. Health care worker groups", S["h2"]),
           _table(S, [["National group", "Includes"],
                      ["Nurse/Midwife", "Registered and enrolled nurses and midwives"],
                      ["Medical", "Medical officers of every grade"],
                      ["Allied health", "Physiotherapists, occupational therapists, dietitians, pharmacists, "
                                        "radiographers and other allied health staff"],
                      ["Personal care staff", "Assistants in nursing and patient care assistants"],
                      ["Support services", "Ward support, cleaning, catering and transport staff"],
                      ["Student", "Nursing, medical and allied health students on placement"]], [4.0, 12.8]),
           Paragraph("5. Hand rub and hand soap use", S["h2"]),
           Paragraph("Alcohol-based hand rub use is reported to the program's quarterly review as the litres of hand "
                     "rub put out on a ward per 1,000 occupied bed days. The litres are those of the bottles the "
                     "store issues to the ward's cost centre (500 mL pump bottles and cartons of 24 personal 100 mL "
                     "bottles), less bottles returned to the store, and of the 1 litre refills the cleaning "
                     "contractor fits to the ward's wall dispensers on its dispenser rounds. Liquid hand soap use is "
                     "reported the same way: the litres of the 500 mL soap pump bottles the store issues to the "
                     "ward's cost centre, less returns, and of the 800 mL cartridges the contractor fits to the soap "
                     "dispensers at the ward's sinks. Occupied bed days are the midnight census's occupied beds "
                     "summed over the days of the period.", S["body"]),
           Paragraph("6. Reporting", S["h2"]),
           Paragraph("Ward results are reported each quarter to the ward's nurse unit manager and to the hospital's "
                     "infection prevention committee, and the hospital's results to the national program three "
                     "times a year.", S["body"])]
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ------------------------------------------------------------------------------------------------ surveillance protocol
def write_definitions(path):
    S = styles("Sans", NAVY)
    doc, fr = _doc(path, "Hospital-onset bloodstream infection surveillance protocol", K.PEOPLE["ipc"],
                   "Infection prevention and control", "%s, Infection Prevention and Control" % K.HOSPITAL,
                   "HOBSI surveillance protocol, version 3")
    st = [Paragraph("Hospital-onset bloodstream infection surveillance protocol", S["h1"]),
          Paragraph("Version 3, effective 1 July 2025. Infection Prevention and Control Service, %s. Replaces "
                    "version 2 (2022)." % K.HOSPITAL, S["meta"]),
          _clause(S, "1", "Scope. Inpatients of the hospital's wards. Patients seen in the emergency department, "
                          "clinics, day units and the renal dialysis unit who are not admitted are outside "
                          "surveillance."),
          _clause(S, "2", "Positive blood culture. A recognised pathogen grown from at least one blood culture set; "
                          "or a common commensal grown from two or more blood culture sets collected within 48 hours "
                          "of each other. A common commensal grown from a single set is a contaminant and is not an "
                          "infection."),
          _clause(S, "3", "Hospital onset. An episode is hospital onset when its first positive blood culture was "
                          "collected more than 48 hours after the patient was admitted to the hospital. A culture "
                          "collected before admission or within 48 hours of it is community onset."),
          _clause(S, "4", "Episodes. A further positive blood culture with the same organism collected less than 14 "
                          "days after the episode's first positive culture belongs to that episode. The same "
                          "organism 14 days or more after it begins a new episode."),
          _clause(S, "5", "Attribution. An episode is attributed to the ward the patient was on when its first "
                          "positive blood culture was collected."),
          _clause(S, "6", "Rate. Hospital-onset bloodstream infections per 10,000 occupied bed days, the occupied bed "
                          "days from the patient administration system's midnight census."),
          Paragraph("Common commensals", S["h2"]),
          Paragraph("Coagulase-negative staphylococci (including <i>Staphylococcus epidermidis</i> and "
                    "<i>Staphylococcus hominis</i>), <i>Micrococcus</i> species, <i>Corynebacterium</i> species "
                    "(diphtheroids), <i>Cutibacterium</i> species, <i>Bacillus</i> species other than <i>B. "
                    "anthracis</i>, and viridans group streptococci.", S["body"])]
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ------------------------------------------------------------------------------------------------ reference workbook
def _style_sheet(sh, widths=None, fill="1D5A57"):
    from openpyxl.styles import Alignment, Font, PatternFill
    for cell in sh[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sh.freeze_panes = "B2"
    for col in sh.columns:
        letter_ = col[0].column_letter
        if widths and letter_ in widths:
            sh.column_dimensions[letter_].width = widths[letter_]
        else:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:200])
            sh.column_dimensions[letter_].width = min(46, max(8, width + 2))


def write_wards(path, side):
    import openpyxl
    floor_of = {u: fl for fl, units in SD.FLOORS.items() for u in units}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Wards"
    ws.append(["code", "ward", "specialty", "beds", "floor", "cost_centre", "audit_application_department",
               "laboratory_location", "notes"])
    for code, name, spec, beds, *_ in K.WARDS:
        note = None
        if code == K.MOVED[0]:
            note = "Moved to the surgical block on 2 February 2026; code S1A from that date."
        ws.append([code, name, spec, beds, "Level " + floor_of[code] if floor_of[code] != "G" else "Ground",
                   R.COST_CENTRE[code][0], R.APP_DEPT[code], R.LAB_LOC[code], note])
        if code == K.MOVED[0]:
            ws.append([K.MOVED[1], "Ward S1A", spec, beds, "Surgical Block Level 1", R.COST_CENTRE[code][0],
                       R.APP_DEPT[K.MOVED[1]], R.LAB_LOC[K.MOVED[1]], "Ward 6A from 2 February 2026."])
    _style_sheet(ws, {"I": 60})
    wa = wb.create_sheet("Areas")
    wa.append(["code", "area", "floor", "audit_application_department", "laboratory_location"])
    for code, name in K.AREAS:
        fl = floor_of[code]
        wa.append([code, name, {"G": "Ground", "S1": "Surgical Block Level 1"}.get(fl, "Level " + fl),
                   R.APP_DEPT[code], R.LAB_LOC[code]])
    _style_sheet(wa)
    wd = wb.create_sheet("Dispensers")
    wd.append(["dispenser_id", "ward_or_area", "floor", "type", "installed", "removed"])
    for did, u, fl, f, t in side.dispensers:
        wd.append([did, u, {"G": "Ground", "S1": "Surgical Block Level 1"}.get(fl, "Level " + fl),
                   side.disp_type[did], f.strftime("%d/%m/%Y"), t.strftime("%d/%m/%Y") if t else None])
    _style_sheet(wd)
    wb.save(path)

