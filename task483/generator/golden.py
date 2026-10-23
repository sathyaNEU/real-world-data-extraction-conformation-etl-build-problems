"""The golden: every owed figure computed from the cut pack (bundle.py), written to generator/figures.json, and the two
deliverables, larkspur_workforce_value_2027-28.xlsx and larkspur_program_tests_2027-28.svg."""
import io
import json
import os

import numpy as np

import bundle as B

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = "larkspur_workforce_value_2027-28.xlsx"
SVG = "larkspur_program_tests_2027-28.svg"


def compute(target):
    I = B.Inputs(target)
    c = B.call(I)
    d0 = B.call(I, "naive", "quarter")
    d_deg = B.call(I, "golden", "quarter")
    d_short = B.call(I, "naive", "golden")
    a = B.ask_a(I)
    b = B.ask_b(I)
    programs = {}
    for p in I.degree:
        x = c["deg"][p]
        programs[p] = {"test": "earnings", "tested": x["tested"], "completers_2024_25": I.c2425.get(p, 0),
                       "with_earnings": x["n"],
                       "figure_thousands": None if not x["tested"] else B.hround(x["figure"] / 1000.0),
                       "passes": x["passes"], "loan_median_usd": b[p][1], "borrowers": b[p][0]}
    for p in I.short:
        x = c["short"][p]
        programs[p] = {"test": "placement", "tested": x["tested"], "completers_2024_25": I.c2425.get(p, 0),
                       "grant_completers": x["n"],
                       "figure_pct": None if not x["tested"] else B.hround(x["figure"]),
                       "working": None if not x["tested"] else int(round(x["working"])),
                       "passes": x["passes"], "tuition_fy2025_usd": a[p]}
    fail = sorted(p for p, v in programs.items() if v["tested"] and not v["passes"])
    ma, cdl = "Medical Assisting", "Commercial Driver License, Class A"
    fig = {
        "call": {"renews": c["renews"], "share_pct": B.hround(c["share"]), "in_passing": c["num"], "tested": c["den"],
                 "failing": fail, "line_pct": 75},
        "programs": programs,
        "components": {
            "medical_assisting": {"pooled_thousands": programs[ma]["figure_thousands"],
                                  "averaged_thousands": B.hround(d0["deg"][ma]["figure"] / 1000.0),
                                  "completers": I.c2425[ma]},
            "cdl": {"dated_pct": programs[cdl]["figure_pct"], "quarter_pct": B.hround(d0["short"][cdl]["figure"]),
                    "completers": I.c2425[cdl], "grant_completers": programs[cdl]["grant_completers"],
                    "working": programs[cdl]["working"]},
            "default_share_pct": B.hround(d0["share"]), "default_in_passing": d0["num"],
            "degree_half_alone_pct": B.hround(d_deg["share"]), "short_half_alone_pct": B.hround(d_short["share"]),
            "untested": sorted(p for p, v in programs.items() if not v["tested"]),
        },
        "stated": {"benchmark_usd": 34200, "benchmark_thousands": 34.2, "placement_line_pct": 70, "share_line_pct": 75, "cpi_factor": 1.0294,
                   "min_with_earnings": 30, "min_grant_completers": 20, "day": 180, "fee_pct": 1.057,
                   "award_years": ["2020-21", "2021-22"], "earnings_years": [2024, 2025]},
        "_unrounded": {"share": c["share"], "default": d0["share"], "degree_half_alone": d_deg["share"],
                       "short_half_alone": d_short["share"],
                       "figures": {p: (c["deg"][p]["figure"] if p in c["deg"] else c["short"][p]["figure"])
                                   for p in programs},
                       "loan_median": {p: b[p][2] for p in I.degree}},
    }
    return fig


def write_xlsx(path, f):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = openpyxl.Workbook()
    bold = Font(name="Calibri", size=11, bold=True)
    thin = Side(style="thin", color="808080")
    fill = PatternFill("solid", fgColor="FCE4D6")
    pr = f["programs"]
    cl = f["call"]

    ws = wb.active
    ws.title = "Designation 2027-28"
    verdict = "renews" if cl["renews"] else "lapses"
    ws["A1"] = "Larkspur Basin Community College: workforce value designation %s for 2027-28" % verdict
    ws["A1"].font = Font(name="Calibri", size=14, bold=True)
    ws["A2"] = ("%d per cent of the 2024-25 completers of tested programs are in passing programs (%s of %s); the "
                "standard needs 75 per cent." % (cl["share_pct"], "{:,}".format(cl["in_passing"]),
                                                 "{:,}".format(cl["tested"])))
    ws["A2"].font = Font(name="Calibri", size=10)
    ws["A3"] = "Grant programs office, Workforce Value Designation Standard sections 3 to 5. Failing programs shaded."
    ws["A3"].font = Font(name="Calibri", size=9, italic=True, color="555555")

    def block(title, head, rows, fmts):
        ws.append([])
        ws.append([title])
        ws.cell(row=ws.max_row, column=1).font = Font(name="Calibri", size=11, bold=True)
        ws.append(head)
        for c in ws[ws.max_row]:
            c.font = bold
            c.alignment = Alignment(wrap_text=True, vertical="bottom")
            c.border = Border(bottom=thin)
        ws.row_dimensions[ws.max_row].height = 44
        for r, shade in rows:
            ws.append(r)
            for i, f in fmts.items():
                ws.cell(row=ws.max_row, column=i + 1).number_format = f
            if shade:
                for c in ws[ws.max_row]:
                    c.fill = fill
    deg = [p for p in pr if pr[p]["test"] == "earnings"]
    sht = [p for p in pr if pr[p]["test"] == "placement"]
    rows = []
    for p in deg:
        x = pr[p]
        res = "Not tested" if not x["tested"] else ("Passes" if x["passes"] else "Fails")
        rows.append(([p, x["with_earnings"] or None, x["figure_thousands"], 34.2 if x["tested"] else None, res,
                      x["completers_2024_25"]], x["tested"] and not x["passes"]))
    block("Credit programs: earnings test", ["Program", "Completers with earnings, 2020-21 and 2021-22 cells",
                                             "Median earnings, thousand 2025 dollars", "Benchmark, thousand dollars",
                                             "Result", "2024-25 completers"], rows,
          {1: "#,##0", 2: "#,##0", 3: "0.0", 5: "#,##0"})
    rows = []
    for p in sht:
        x = pr[p]
        res = "Not tested" if not x["tested"] else ("Passes" if x["passes"] else "Fails")
        rows.append(([p, x["grant_completers"], x["working"], x["figure_pct"], 70 if x["tested"] else None, res,
                      x["completers_2024_25"]], x["tested"] and not x["passes"]))
    block("Short-term programs: placement test", ["Program", "Grant completers, sections ending July 2024 to June 2025",
                                                  "Working on day 180", "Share working, per cent", "Line, per cent",
                                                  "Result", "2024-25 completers"], rows,
          {1: "#,##0", 2: "#,##0", 3: "0", 4: "0", 6: "#,##0"})
    ws.append([])
    for label, v, fmt in (("Tested programs' 2024-25 completers", cl["tested"], "#,##0"),
                          ("In passing programs", cl["in_passing"], "#,##0"),
                          ("Share in passing programs, per cent", cl["share_pct"], "0")):
        ws.append([label] + [None] * 5 + [v])
        ws.cell(row=ws.max_row, column=1).font = bold
        ws.cell(row=ws.max_row, column=7).font = bold
        ws.cell(row=ws.max_row, column=7).number_format = fmt
    for col, wd in zip("ABCDEFG", (36, 16, 15, 14, 12, 11, 12)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "B5"

    s2 = wb.create_sheet("Grant tuition FY2025")
    s2.append(["Short-term program", "Grant tuition and fees paid, July 2024 to June 2025 (USD)"])
    for c in s2[1]:
        c.font = bold
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
        c.border = Border(bottom=thin)
    for p in [p for p in pr if pr[p]["test"] == "placement"]:
        s2.append([p, pr[p]["tuition_fy2025_usd"]])
    s2.append(["All short-term programs", sum(pr[p]["tuition_fy2025_usd"] for p in pr if pr[p]["test"] == "placement")])
    for c in s2[s2.max_row]:
        c.font = bold
        c.border = Border(top=thin)
    for row in s2.iter_rows(min_row=2, max_row=s2.max_row):
        row[1].number_format = "#,##0"
    s2.column_dimensions["A"].width = 36
    s2.column_dimensions["B"].width = 24
    s2.row_dimensions[1].height = 46
    s2.freeze_panes = "A2"

    s3 = wb.create_sheet("Loan principal")
    s3.append(["Program", "Median Direct Loan principal, 2024-25 completers who borrowed (USD)"])
    for c in s3[1]:
        c.font = bold
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
        c.border = Border(bottom=thin)
    for p in [p for p in pr if pr[p]["test"] == "earnings"]:
        s3.append([p, pr[p]["loan_median_usd"]])
    for row in s3.iter_rows(min_row=2, max_row=s3.max_row):
        row[1].number_format = "#,##0"
    s3.column_dimensions["A"].width = 30
    s3.column_dimensions["B"].width = 26
    s3.row_dimensions[1].height = 46
    s3.freeze_panes = "A2"

    s4 = wb.create_sheet("Notes")
    notes = [
        ("Median earnings", "The release publishes each campus's cohort as a cell of completers with earnings and five "
         "percentiles. Each cell's earnings distribution is rebuilt by joining its percentiles straight (2024 earnings "
         "times 1.0294), and the program's median is the point where the cells' pooled share of completers reaches "
         "half, weighting each cell by its completers with earnings. Paralegal Studies has no published cell."),
        ("Working on day 180", "Day 180 is the section's last day plus 180 days. A grant completer counts when a job "
         "paid in the quarter holding day 180 was running on it: paid in the quarters before and after as well, or a "
         "start or an exit inside the quarter dated from that job's pay there against its pay in its adjacent full "
         "quarter. Amended wage reports replace the originals. Forklift and Material Handling has 18 grant completers "
         "and is not tested."),
        ("Grant tuition", "The board's payments dated July 2024 to June 2025, each split by the college's invoice lines: "
         "tuition and the section fee to the section's program, each refund credit to the program of the participant "
         "it refunds."),
        ("Loan principal", "Each completer's Direct Subsidized and Unsubsidized disbursements at the college, returns "
         "netted and parent PLUS loans left out. Disbursements from 1 July 2023 are exported net of the 1.057 per cent "
         "origination fee and are grossed back to the amount borrowed."),
        ("Rounding", "Earnings to whole thousands of dollars, shares to whole per cent, tuition to whole dollars and loan "
         "principal to the nearest 10 dollars, each from the unrounded figure; each figure is compared with its line "
         "before rounding."),
    ]
    s4.append(["Item", "Note"])
    for c in s4[1]:
        c.font = bold
    for k, v in notes:
        s4.append([k, v])
    for row in s4.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[0].alignment = Alignment(vertical="top")
    s4.column_dimensions["A"].width = 20
    s4.column_dimensions["B"].width = 100
    wb.save(path)
    import postprocess as PP
    PP.normalize_ooxml(path, "Grant Programs Office", "Grant Programs Office", "2026-10-27T15:12:00Z",
                       "2026-11-06T22:04:31Z", excel=True, company="State Workforce Development Board")


def write_svg(path, f):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Patch
    matplotlib.rcParams["svg.hashsalt"] = "task483-larkspur"
    matplotlib.rcParams["svg.fonttype"] = "none"
    matplotlib.rcParams["font.family"] = "DejaVu Sans"
    pr = f["programs"]
    un = f["_unrounded"]["figures"]
    pass_c, fail_c, ink, muted, line = "#2a6fb0", "#c8641a", "#1f2328", "#59636e", "#3d434a"
    deg = sorted([p for p in pr if pr[p]["test"] == "earnings" and pr[p]["tested"]], key=lambda p: -un[p])
    sht = sorted([p for p in pr if pr[p]["test"] == "placement" and pr[p]["tested"]], key=lambda p: -un[p])
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(11.6, 6.4), dpi=100)
    for ax, names, scale, ln, lab, unit in ((a1, deg, 1000.0, 34.2, "34.2 thousand", "thousand dollars"),
                                            (a2, sht, 1.0, 70.0, "70 per cent", "per cent")):
        ys = np.arange(len(names))[::-1]
        vals = [un[p] / scale for p in names]
        cols = [pass_c if pr[p]["passes"] else fail_c for p in names]
        ax.barh(ys, vals, height=0.64, color=cols, edgecolor="white", linewidth=1.5, zorder=3)
        for y, p, v in zip(ys, names, vals):
            shown = pr[p]["figure_thousands"] if unit.startswith("thousand") else pr[p]["figure_pct"]
            ax.text(v - max(vals) * 0.015, y, "%d" % shown, ha="right", va="center", fontsize=8.5, color="white",
                    fontweight="bold", zorder=5)
        ax.axvline(ln, color=line, linewidth=1.3, linestyle=(0, (4, 3)), zorder=4)
        ax.text(ln, len(names) - 0.35, " " + lab, color=line, fontsize=8, va="bottom", ha="left")
        ax.set_yticks(ys)
        ax.set_yticklabels(names, fontsize=8.5, color=ink)
        ax.set_xlim(0, max(vals) * 1.14)
        ax.set_ylim(-0.7, len(names) + 0.3)
        ax.tick_params(axis="x", labelsize=8, colors=muted)
        ax.grid(axis="x", color="#e6e8eb", linewidth=0.8, zorder=0)
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
        ax.spines["left"].set_color("#9aa1a9")
        ax.spines["bottom"].set_color("#9aa1a9")
    a1.set_title("Credit programs: median earnings in the third year\nafter the award, 2025 thousands of dollars",
                 loc="left", fontsize=9, color=muted)
    a2.set_title("Short-term programs: grant completers working\non day 180, per cent", loc="left", fontsize=9,
                 color=muted)
    cl = f["call"]
    fig.legend(handles=[Patch(color=pass_c, label="Passes its test"), Patch(color=fail_c, label="Fails its test")],
               loc="upper right", frameon=False, fontsize=8.5, ncol=2, bbox_to_anchor=(0.985, 0.965))
    fig.suptitle("Workforce value designation %s for 2027-28: %d per cent of completers in passing programs, "
                 "75 needed" % ("renews" if cl["renews"] else "lapses", cl["share_pct"]),
                 x=0.012, ha="left", fontsize=11.5, fontweight="bold", color=ink)
    fig.text(0.012, 0.012, "Larkspur Basin Community College. Earnings: Postsecondary Earnings Outcomes, 2024 and 2025 "
             "editions. Placement: Career Training Grant roster and wage records.", fontsize=7.5, color=muted)
    fig.subplots_adjust(left=0.2, right=0.985, top=0.83, bottom=0.08, wspace=0.62)
    buf = io.StringIO()
    fig.savefig(buf, format="svg", metadata={"Date": None, "Creator": "Grant Programs Office"})
    plt.close(fig)
    svg = buf.getvalue().replace('id="matplotlib.', 'id="').replace("'DejaVu Sans'",
                                                                   "'DejaVu Sans', Arial, Helvetica, sans-serif")
    assert "matplotlib" not in svg.lower()
    open(path, "w", encoding="utf-8", newline="\n").write(svg)


def main(target=None):
    target = target or os.path.join(os.path.dirname(HERE), "target")
    f = compute(target)
    json.dump(f, open(os.path.join(HERE, "figures.json"), "w"), indent=1, sort_keys=True)
    out = os.path.join(os.path.dirname(HERE), "golden")
    os.makedirs(out, exist_ok=True)
    for fn in os.listdir(out):
        os.remove(os.path.join(out, fn))
    write_xlsx(os.path.join(out, XLSX), f)
    write_svg(os.path.join(out, SVG), f)
    return f


if __name__ == "__main__":
    import sys
    f = main(sys.argv[1] if len(sys.argv) > 1 else None)
    print(json.dumps({k: v for k, v in f.items() if not k.startswith("_")}, indent=1)[:5000])
