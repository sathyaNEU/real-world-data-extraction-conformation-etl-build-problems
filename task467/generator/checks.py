"""The generator's assertions for task467: the three rules executed on the clean truth (R1 to R3c), every row of the
semantic fork table (SF1 to SF13), the ladder and its margins, the graded figures' distance to their rounding edges,
the Layer 1 items' neutrality, the asks' devices and stops, the separation of the call from the asks, the distractors,
the goldens against figures.json, the input gates and the realism sweep.

    python3 task467/generator/checks.py      (after build.py has written target/ and golden/)
"""
import csv
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import numpy as np                # noqa: E402
import pandas as pd               # noqa: E402
import pyarrow.parquet as pq      # noqa: E402

import bundle as B                # noqa: E402
import golden as GO               # noqa: E402
import methods as M               # noqa: E402
import sim as S                   # noqa: E402
import world as W                 # noqa: E402

TASK = os.path.dirname(HERE)
PERSONAS = {"Olivia White", "Joseph Heath", "Johnny Houston"}      # guard.py names --geo Canada --seed 467
D = dt.date.fromisoformat
POST26 = (D("2025-01-01"), D("2026-08-31"))                      # 2025 with 2026 to date
M25 = ["2025-%02d" % m for m in range(1, 13)]
M26 = ["2026-%02d" % m for m in range(1, 9)]
DASH = re.compile("[–—]")


# ---------------------------------------------------------------- helpers
def kept(c, zs, pre, post):
    return (M.n(c, post, "P", zs) / M.n(c, pre, "P", zs)) / (M.n(c, post, "D", zs) / M.n(c, pre, "D", zs))


def solve(c, amount, groups=None, pre=M.PRE, post=M.BASE, base=M.BASE):
    """The base year's pickups restated by each group's kept share over (pre, post), the share and the fee."""
    groups = groups or M.strata()
    rec = M.n(c, base, "P")
    q = sum(M.n(c, base, "P", zs) / kept(c, zs, pre, post) for zs in groups.values())
    return dict(rec=rec, q=q, q1000=M.round_to(q, 1000), share=100 * (1 - rec / q),
                pct=M.round_to(100 * (1 - rec / q), 1), fee=M.fee_for(amount, q), cq=amount / q)


def merge(parts):
    c = pd.concat(parts, ignore_index=True)
    return c.groupby(["date", "dir", "zone"], as_index=False)["n"].sum()


def tcounts(t):
    return M.counts_from_trips(B.as_trips(t))


def arrays_counts(ep, ed):
    rows = []
    for direction, arr in (("P", ep), ("D", ed)):
        di, zj = np.nonzero(arr)
        rows.append(pd.DataFrame({"date": S.DAY64[di], "dir": direction, "zone": np.array(S.ZONES)[zj],
                                  "n": arr[di, zj]}))
    return pd.concat(rows, ignore_index=True)


def pdf_text(path):
    out = subprocess.run(["pdftotext", "-raw", path, "-"], capture_output=True, text=True, check=True).stdout
    return re.sub(r"\s+", " ", out)


def doc_texts(target):
    out = {}
    for f in sorted(os.listdir(target)):
        p = os.path.join(target, f)
        if f.endswith(".pdf"):
            out[f] = pdf_text(p)
        elif f.endswith(".xlsx"):
            import openpyxl
            wb = openpyxl.load_workbook(p)
            out[f] = " ".join(str(v) for ws in wb.worksheets for r in ws.iter_rows(values_only=True) for v in r
                              if v is not None)
    return out


def docx_text(path):
    from docx import Document
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            parts.append(" | ".join(c.text for c in r.cells))
    return "\n".join(parts)


def ledger_of(meter_on, brokerage):
    return M.ledger(pd.DataFrame({"meter_on": meter_on, "brokerage": np.asarray(brokerage).astype(int)})).set_index(
        ["month", "brokerage"])


def misses(alt, gold, months, col="pickups"):
    g = gold[[m in months for m in gold.index.get_level_values(0)]]
    a = alt.reindex(g.index) if isinstance(alt, pd.DataFrame) else alt.reindex(g.index).to_frame(col)
    return int((a[col].fillna(-1).round(2) != g[col].round(2)).sum()), len(g)


def copy_without(target, keys):
    tmp = tempfile.mkdtemp(dir=TASK, suffix=".chk")
    drop = {W.FILES[k] for k in keys}
    for f in os.listdir(target):
        if f not in drop:
            shutil.copy2(os.path.join(target, f), os.path.join(tmp, f))
    return tmp


def passengers(target):
    import openpyxl
    ws = openpyxl.load_workbook(B.path(target, "passengers"), data_only=True).active
    out = {}
    for r in ws.iter_rows(values_only=True):
        if r and isinstance(r[0], str) and re.fullmatch(r"\d{4}-\d{2}", r[0]):
            out[r[0]] = (int(r[1]), int(r[2]))
    return out


def run(target, golden, x=None, meta_path=None, fig_path=None):
    bad = []

    def check(label, cond, detail=""):
        print("%s  %s%s" % ("PASS" if cond else "FAIL", label, ("  [" + str(detail) + "]") if detail != "" else ""))
        if not cond:
            bad.append(label)

    x = x or S.everything()
    F = json.load(open(fig_path or os.path.join(HERE, "figures.json"), encoding="utf-8"))
    meta = json.load(open(meta_path or os.path.join(TASK, "metadata.json"), encoding="utf-8"))
    files = sorted(os.listdir(target))
    docs = doc_texts(target)
    amount = B.budget_amount(target)

    # the bundle, read once
    t24raw = B._read(B.path(target, "trips2024"))
    t24 = B.drop_voids(t24raw)
    t25 = B._read(B.path(target, "trips2025"))
    arch = B.archive(target)
    d26raw = pd.concat([B._read(B.path(target, k)).assign(drop=k) for k in W.DROP_KEYS], ignore_index=True)
    d26nv = B.drop_voids(d26raw)
    d26 = B.drop_copies(d26nv)
    c = merge([arch, tcounts(t24), tcounts(t25), tcounts(d26)])
    g = solve(c, amount)
    q_drop = M.restated(c)
    fx = GO.compute(target)

    # ------------------------------------------------------------------ the call
    check("the call: $%.2f a pickup; 2025 recorded %d, restated %.1f (%d to the thousand), the rank's share %.2f per "
          "cent, the amount $%.2f" % (g["fee"], g["rec"], g["q"], g["q1000"], g["share"], amount),
          g["fee"] == 6.50 and F["call"]["fee_2027"] == "6.50" and fx["fee"] == 6.50 and amount == W.AMOUNT_2027)
    check("the critical components in figures.json recompute: %s recorded, %s restated, %s per cent" % (
        F["components"]["recorded_2025"], F["components"]["restated_2025"], F["components"]["share_cost_pct"]),
          F["components"]["recorded_2025"] == "{:,}".format(g["rec"]) and
          F["components"]["restated_2025"] == "{:,}".format(g["q1000"]) and F["components"]["share_cost_pct"] == "%d" % g["pct"]
          and F["components"]["amount_2027"] == "{:,}".format(int(amount)))

    # ------------------------------------------------------------------ R1: the field's raw material ships
    zones_seen = {d_: set(c[(c["dir"] == d_) & (c["date"] >= np.datetime64(M.PRE[0])) &
                            (c["date"] <= np.datetime64(M.BASE[1]))]["zone"]) for d_ in "PD"}
    check("R1 the pack ships the series' raw material: trip records by meter start, pickup zone and drop-off zone for "
          "2024 to August 2026 and the 2023 daily archive, both directions in all %d city zones" % len(W.CITY_ZONES),
          zones_seen["P"] == set(W.CITY_ZONES) == zones_seen["D"] and len(t24) > 100000 and len(t25) > 100000)
    cols = {}
    for f in files:
        p = os.path.join(target, f)
        if f.endswith(".parquet"):
            cols[f] = pq.ParquetFile(p).schema_arrow.names
        elif f.endswith(".csv"):
            cols[f] = open(p, encoding="utf-8").readline().strip().split(",")
    flagged = {f: [k for k in v if re.search(r"station|line|rank|restat|kept|share|effect(?!ive)|counterfact|adjust", k, re.I)]
               for f, v in cols.items()}
    check("R1 no shipped column marks a station zone, a restated count or the rank's share (the zone list carries "
          "zone_id, zone_name and district)", not any(flagged.values()) and
          cols[W.FILES["zones"]] == ["zone_id", "zone_name", "district"], {f: v for f, v in flagged.items() if v})
    rung_q = {"R0": M.rung0(c), "R1": M.rung1(c), "R2": g["q"], "none": float(g["rec"])}
    tokens = set()
    for v in rung_q.values():
        tokens |= {"{:,}".format(int(round(v))), "{:,}".format(M.round_to(v, 1000)), str(int(round(v)))}
    hits = {f: sorted(t for t in tokens if re.search(r"(?<![\d,.])%s(?![\d,])" % re.escape(t), txt))
            for f, txt in docs.items()}
    check("R1 no shipped document or workbook states 2025's recorded or restated pickups under any rung (%s)" %
          ", ".join("%s %.0f" % kv for kv in rung_q.items()), not any(hits.values()), {f: v for f, v in hits.items() if v})
    check("R1 the decisive computation reads zone-level counts: the golden restatement equals the sum over the two "
          "zone groups of 2025 drop-offs at the group's pre-move ratio (%.1f)" % q_drop, abs(q_drop - g["q"]) < 1e-6)

    # ------------------------------------------------------------------ R3: the analyst test on the clean truth
    tr = S.truth(x)
    ct = M.counts_from_trips(tr)
    ct = ct[ct["date"] <= np.datetime64(W.LAST_DAY)]
    t_r0 = M.fee_for(amount, M.rung0(ct))
    t_r1 = M.fee_for(amount, M.rung1(ct))
    t_r2 = M.fee_for(amount, M.restated(ct))
    ce = arrays_counts(x["ep"], x["ed"])
    e_r0, e_r1, e_r2 = (M.fee_for(amount, v) for v in (M.rung0(ce), M.rung1(ce), M.restated(ce)))
    check("R3a the analyst test on the clean truth: every field read right and every filed rule applied (the by-law's "
          "fee-liable pickup, FP-07's base year, its taxi loading arrangement, no growth, the lowest 25-cent step), "
          "the rank's effect read as the pickups' own fall files $%.2f (and $%.2f on the noise-free expectation)" % (
              t_r0, e_r0), t_r0 == 5.75 and e_r0 == 5.75)
    check("R3b the decisive step alone (pickups against drop-offs within the line's station zones and within the rest) "
          "files the golden $%.2f on the clean truth (noise-free $%.2f)" % (t_r2, e_r2), t_r2 == 6.50 and e_r2 == 6.50)
    check("R3b the lower rung's single correction (all pickups against all drop-offs) stays wrong on the clean truth: "
          "$%.2f (noise-free $%.2f)" % (t_r1, e_r1), t_r1 == 7.00 and e_r1 == 7.00)
    kt = {k: 100 * (1 - kept(ce, zs, M.PRE, M.BASE)) for k, zs in M.strata().items()}
    check("R3b the clean truth's generating effect: the rank keeps %.3f of pickups in every zone; the noise-free "
          "within-group shares are %.2f and %.2f per cent" % (W.RANK, kt["station"], kt["other"]),
          all(abs(v - 100 * (1 - W.RANK)) < 0.3 for v in kt.values()))
    monthly = c.assign(m=pd.to_datetime(c["date"]).dt.strftime("%Y-%m")).groupby(["m", "dir"])["n"].sum().unstack()
    pairs = list(zip(monthly["P"], monthly["D"]))
    zm = c.assign(m=pd.to_datetime(c["date"]).dt.strftime("%Y-%m")).pivot_table(index="m", columns=["dir", "zone"],
                                                                                 values="n", aggfunc="sum",
                                                                                 fill_value=0)
    daily = c[c["dir"] == "P"].groupby("date")["n"].sum().to_numpy()
    runs = {tuple(daily[i:i + 14]) for i in range(len(daily) - 13)}
    check("R3c no copied, held or swapped block: %d months with distinct pickup and drop-off totals, no zone-month "
          "vector repeated, no 14-day run of daily pickups repeated" % len(pairs),
          len(set(pairs)) == len(pairs) and len(zm.drop_duplicates()) == len(zm) and len(runs) == len(daily) - 13)
    money = {f: sorted(set(re.findall(r"\$(?:6\.25|6\.50|7\.00|7\.25)\b", t))) for f, t in docs.items()}
    check("R3c no document names a candidate 2027 fee or the decisive step (no $6.25, $6.50, $7.00 or $7.25; no "
          "station-zone, stratum or comparison language)", not any(money.values()) and not any(
              re.search(r"stratum|strata|comparison group|station zone|Simpson|difference-in-differences", t, re.I)
              for t in docs.values()), {f: v for f, v in money.items() if v})
    vol = {f: re.findall(r"\d{1,3}(?:,\d{3})+ (?:fee-liable )?(?:pickups|trips)", t) for f, t in docs.items()}
    check("R3c no memo, draft paper or stakeholder figure: no file is a memo, briefing or note, and no document quotes "
          "a pickup or trip volume", not any(re.search(r"memo|brief|note_|draft|analysis", f) for f in files)
          and not any(vol.values()), {f: v for f, v in vol.items() if v})
    sched = docs[W.FILES["bylaw"]]
    check("R3c Schedule A sets no 2027 fee and the budget states no volume or fee", "2027 To be set by the Board under "
          "section 4" in sched and not re.search(r"pickups|per pickup|\$\d\.\d\d", docs[W.FILES["budget"]]))

    # ------------------------------------------------------------------ R2: every Layer 1 item real and neutral
    def same(r):
        return r["fee"] == 6.50 and r["q1000"] == g["q1000"] and r["pct"] == g["pct"] and r["rec"] == g["rec"]

    raw24 = pq.read_table(B.path(target, "trips2024")).to_pandas()
    raw25 = pq.read_table(B.path(target, "trips2025")).to_pandas()
    pos = raw24.copy()
    pos.columns = list(raw25.columns)
    pos = B.drop_voids(pos)
    r_pos = solve(merge([arch, tcounts(pos), tcounts(t25)]), amount)
    check("R2 the version 2 and version 3 field names (2024 and 2025 annual files): a union by position gives the same "
          "rows and the same call ($%.2f)" % r_pos["fee"], list(B.V2_TO_V3.get(k, k) for k in raw24.columns) ==
          list(raw25.columns) and same(r_pos))
    r_vk = solve(merge([arch, tcounts(t24raw), tcounts(t25)]), amount)
    r_vo = solve(merge([arch, tcounts(t24raw[t24raw["total"] >= 0]), tcounts(t25)]), amount)
    nvoid = int((t24raw["total"] < 0).sum())
    check("R2 the 2024 file's %d voided trips with their reversal rows: counted as trips (%.0f, $%.2f) or with only the "
          "reversals dropped (%.0f, $%.2f), the call and every component stay" % (nvoid, r_vk["q"], r_vk["fee"],
                                                                                  r_vo["q"], r_vo["fee"]),
          nvoid > 300 and same(r_vk) and same(r_vo))
    sfy = {y: float((t["sf_flag"] == "Y").mean()) for y, t in ((2024, t24), (2025, t25))}
    sf = merge([arch, tcounts(t24[t24["sf_flag"] == "N"]), tcounts(t25[t25["sf_flag"] == "N"])])
    r_sf = solve(sf, amount, base=M.BASE)
    check("R2 the store-and-forward records (sf_flag Y, %.1f and %.1f per cent of the 2024 and 2025 rows): the annual "
          "files place them by meter start, and they fall evenly on both directions, so the rank's share without them "
          "is the same %d per cent (%.2f)" % (100 * sfy[2024], 100 * sfy[2025], r_sf["pct"], r_sf["share"]),
          r_sf["pct"] == g["pct"] and all(0.02 < v < 0.05 for v in sfy.values()))
    r_na = solve(merge([tcounts(t24), tcounts(t25)]), amount, pre=(D("2024-01-01"), M.PRE[1]))
    check("R2 2023 kept only as the daily archive: left out (the pre-move window from 1 January 2024) the call stays "
          "($%.2f, %d)" % (r_na["fee"], r_na["q1000"]), same(r_na))

    def with26(t26):
        return solve(merge([arch, tcounts(t24), tcounts(t25), tcounts(t26)]), amount, post=POST26)

    rawu = pd.concat([pq.read_table(B.path(target, k)).to_pandas().assign(drop=k) for k in W.DROP_KEYS],
                     ignore_index=True)
    rawu = rawu.rename(columns={k: "airport_fee" for k in rawu.columns if k.lower() == "airport_fee"})
    for k in ("meter_on", "meter_off"):
        rawu[k] = rawu[k].dt.tz_localize(None)
    r26 = with26(d26)
    r_utc = with26(B.drop_copies(B.drop_voids(rawu)))
    r_cp = with26(d26nv)
    r_vk26 = with26(B.drop_copies(d26raw[d26raw["total"] >= 0]))
    dm = d26.copy()
    dm["meter_on"] = pd.to_datetime(["2026-%s-01" % k[-2:] for k in dm["drop"]])
    r_dm = with26(dm)
    check("R2 the 2026 monthly files with 2026 added to the comparison: handled (%.2f per cent, $%.2f), read in UTC "
          "(%.2f, $%.2f), re-sent copies kept (%.2f, $%.2f), voided originals kept (%.2f, $%.2f), placed by file "
          "month (%.2f, $%.2f): the share's whole per cent and the fee stay" % (
              r26["share"], r26["fee"], r_utc["share"], r_utc["fee"], r_cp["share"], r_cp["fee"], r_vk26["share"],
              r_vk26["fee"], r_dm["share"], r_dm["fee"]),
          all(r["fee"] == 6.50 and r["pct"] == g["pct"] for r in (r26, r_utc, r_cp, r_vk26, r_dm)))
    noc = merge([arch, tcounts(t24.drop(columns=["brokerage_id", "airport_fee"])),
                 tcounts(t25.drop(columns=["brokerage_id", "airport_fee"]))])
    check("R2 the fee column (renamed in case in the June 2026 file; January's meters at last year's fee) and the "
          "annual file's brokerage label: the call reads neither, and without them it is unchanged",
          same(solve(noc, amount)) and (t25["airport_fee"][t25["pu_zone"] == 1] == 5.50).sum() > 0)

    # ------------------------------------------------------------------ the semantic fork table
    rec_both = M.n(c, M.BASE, "P") + M.n(c, M.BASE, "D")
    bylaw, fp07 = docs[W.FILES["bylaw"]], docs[W.FILES["policy"]]
    f_both = M.fee_for(amount, rec_both / (1 - g["share"] / 100))
    f_do = M.fee_for(amount, M.restated(c) * M.n(c, M.BASE, "D") / g["rec"])
    check("SF1 fee-liable pickup: every trip touching zone 1 ($%.2f) or the trips ending there ($%.2f) give other fees; "
          "by-law s.2 defines a trip a licensed taxi begins at the airport, voided trips out" % (f_both, f_do),
          f_both != 6.50 and f_do != 6.50 and "fee-liable pickup means a trip a licensed taxi begins at the airport"
          in bylaw and "voided is not a trip" in bylaw)
    tw = (D("2025-09-01"), D("2026-08-31"))
    r_tw = solve(c, amount, base=tw)
    check("SF2 the most recent complete year: the trailing twelve months to August 2026 restate to %.0f and the same "
          "$%.2f; FP-07 4.2 and 1.3 file calendar 2025" % (r_tw["q"], r_tw["fee"]), r_tw["fee"] == 6.50 and
          "most recent complete fiscal year" in fp07 and "fiscal year is the calendar year" in fp07)
    check("SF3 taxi loading arrangement: restating the line away as well gives the pre-move pickups (%.0f) and $%.2f; "
          "FP-07 4.2 defines the arrangement as the rank's and holding lot's location and operation" % (
              M.rung0(c), M.fee_for(amount, M.rung0(c))),
          M.fee_for(amount, M.rung0(c)) == 5.75 and "location and operation of the taxi rank and the taxi holding lot"
          in fp07)
    check("SF4 restated: each group's pickups over its kept share and each group's drop-offs at its pre-move ratio are "
          "one figure (%.4f against %.4f)" % (g["q"], q_drop), abs(g["q"] - q_drop) < 1e-6)
    wins = {"calendar January to October 2024 before": dict(pre=(D("2024-01-01"), D("2024-10-31"))),
            "calendar months November 2023 to October 2024": dict(pre=(D("2023-11-01"), D("2024-10-31"))),
            "all of 2023 and 2024 before the move": dict(pre=(D("2023-01-01"), M.PRE[1])),
            "the twelve months after the move": dict(post=(D("2024-11-04"), D("2025-11-03"))),
            "2025 with 2026 to August": dict(post=POST26),
            "November 2024 to August 2026": dict(post=(D("2024-11-04"), D("2026-08-31"))),
            "January to October 2024 against 2025 to August 2026": dict(pre=(D("2024-01-01"), D("2024-10-31")),
                                                                         post=POST26)}
    wr = {k: solve(c, amount, **v) for k, v in wins.items()}
    check("SF5 measurement windows: %d window pairs restate to %.0f to %.0f, every one %d, %d per cent and $6.50" % (
        len(wr), min(r["q"] for r in wr.values()), max(r["q"] for r in wr.values()), g["q1000"], g["pct"]),
          all(same(r) for r in wr.values()), {k: round(r["q"], 1) for k, r in wr.items()})
    tz = {f: str(pq.ParquetFile(os.path.join(target, f)).schema_arrow.field(
        "meter_on" if "meter_on" in pq.ParquetFile(os.path.join(target, f)).schema_arrow.names else "pickup_datetime").type)
          for f in files if f.endswith(".parquet")}
    check("SF6 day convention: the by-law reckons days, months and years in Velmont local time; the annual files "
          "carry local time without a zone, the 2026 files UTC with it; the share with 2026 added is the same in UTC "
          "and local days", all(("UTC" in v) == ("2026-" in f) for f, v in tz.items()) and r_utc["pct"] == r26["pct"]
          and "Days, months and years are reckoned in Velmont local time" in bylaw, tz)
    near = round(amount / g["q"] / 0.25) * 0.25
    near_v = [round(amount / v / 0.25) * 0.25 for v in [g["q"]] + [r["q"] for r in wr.values()]]
    check("SF7 rounding: the nearest 25 cents gives $%.2f on every window (the same call); by-law s.4 files the lowest "
          "multiple of $0.25 that recovers the amount" % near, set(near_v) == {6.50} and "lowest multiple of $0.25" in bylaw,
          near_v)
    import openpyxl
    lines = {r[0]: (r[1], r[2]) for r in openpyxl.load_workbook(B.path(target, "budget"), data_only=True).active.iter_rows(
        values_only=True) if isinstance(r[0], int)}
    gross = lines[8][1]
    check("SF8 the amount: the program's total cost ($%.2f) gives $%.2f; by-law s.4 files the amount the budget assigns "
          "to the fee, the budget's line 12" % (gross, M.fee_for(gross, g["q"])),
          M.fee_for(gross, g["q"]) > 6.50 and lines[12][0] == "Amount assigned to the taxi pickup fee" and
          "assigns to the taxi pickup fee" in bylaw)
    f_g = M.fee_for(amount, g["q"] * (1 + W.GROWTH) ** 2)
    check("SF9 no traffic growth: grown at the passengers' 3 per cent to 2027 the fee is $%.2f; FP-07 4.2 assumes none"
          % f_g, f_g != 6.50 and "No traffic growth is assumed" in fp07)
    up = M.round_to(100 * (g["q"] / g["rec"] - 1), 1)
    prompt = open(os.path.join(TASK, "prompt.txt"), encoding="utf-8").read()
    check("SF10 the share: lost over restated is %d per cent, the uplift over recorded %d; the prompt asks for the "
          "share of 2025's restated pickups" % (g["pct"], up), g["pct"] == 11 and up == 12 and
          "the share of 2025's restated pickups the temporary rank cost" in prompt)
    note = docs[W.FILES["notice"]]
    check("SF11 the budget year's arrangement: the notice returns the rank to the forecourt on 1 January 2027, so all "
          "of 2027 loads there", "Friday 1 January 2027" in note and "taxi holding lot (P7)" in note)
    check("SF12 the 2024 file's voided trips: the reversal dropped and the original kept, the call stays ($%.2f)" %
          r_vo["fee"], same(r_vo))
    check("SF13 the 2023 archive against trip records: without it the call stays ($%.2f)" % r_na["fee"], same(r_na))

    # ------------------------------------------------------------------ the ladder
    q0, q1 = M.rung0(c), M.rung1(c)
    lad = {"R0": M.fee_for(amount, q0), "R1": M.fee_for(amount, q1), "R2": g["fee"], "none": M.fee_for(amount, g["rec"])}
    check("ladder: R0 the pickups' own fall $%.2f, R1 all pickups against all drop-offs $%.2f, R2 within the line's "
          "station zones and the rest $%.2f (golden), no restatement $%.2f" % tuple(lad.values()),
          lad == {"R0": 5.75, "R1": 7.00, "R2": 6.50, "none": 7.25})
    check("ladder: every correction from R0 walks the fee up, only the decisive move turns it down, and no two rungs "
          "share a fee", lad["R0"] < lad["R1"] > lad["R2"] > lad["R0"] and len(set(lad.values())) == 4)
    shares = {"R0": 100 * (1 - g["rec"] / q0), "R1": 100 * (1 - g["rec"] / q1), "R2": g["share"]}
    check("ladder: the rank's share is %.1f, %.1f and %.1f per cent on the three rungs" % tuple(shares.values()),
          M.round_to(shares["R0"], 1) == 22 and M.round_to(shares["R1"], 1) == 3 and g["pct"] == 11 and
          F["rivals"]["own_fall"]["share_pct"] == "22" and F["rivals"]["all_dropoffs"]["share_pct"] == "3")

    def margin(v, lo, hi):
        return min((v - lo) / lo, (hi - v) / v)

    cqs = {"R0": (amount / q0, 5.50, 5.75), "R1": (amount / q1, 6.75, 7.00), "R2": (g["cq"], 6.25, 6.50),
           "none": (amount / g["rec"], 7.00, 7.25)}
    ms = {k: margin(*v) for k, v in cqs.items()}
    check("bins: the amount over each rung's pickups sits inside its 25-cent step with at least 1.2 per cent clear "
          "(%s)" % ", ".join("%s %.4f, %.1f%%" % (k, v[0], 100 * ms[k]) for k, v in cqs.items()),
          all(m > 0.012 for m in ms.values()) and ms["R2"] > 0.018)
    kept_s = {k: 100 * (1 - kept(c, zs, M.PRE, M.BASE)) for k, zs in M.strata().items()}
    check("strata agree: the rank's share is %.2f per cent in the station zones and %.2f in the rest, both %d" % (
        kept_s["station"], kept_s["other"], g["pct"]), abs(kept_s["station"] - kept_s["other"]) < 0.8 and
          all(M.round_to(v, 1) == 11 for v in kept_s.values()) and F["strata_pct"] == {
              "station_zones": "%.1f" % kept_s["station"], "other_zones": "%.1f" % kept_s["other"]})
    zq = M.restated_by_zone(c)
    zq0 = M.restated_by_zone(c, floor=0)
    check("C1 every zone its own group: %.1f (floor 30) and %.1f (no floor) restated, %d and $%.2f" % (
        zq, zq0, M.round_to(zq, 1000), M.fee_for(amount, zq)),
          M.round_to(zq, 1000) == g["q1000"] == M.round_to(zq0, 1000) and M.fee_for(amount, zq) == 6.50)
    zk = {z: 100 * (1 - kept(c, [z], M.PRE, M.BASE)) for z in W.CITY_ZONES if M.n(c, M.PRE, "P", [z]) > 1500}
    check("C1 every zone with more than 1,500 pre-move pickups shows the rank's share between %.1f and %.1f per cent"
          % (min(zk.values()), max(zk.values())), all(9.5 < v < 12.5 for v in zk.values()) and len(zk) >= 15)
    fall = {k: 100 * (1 - M.n(c, M.BASE, "D", zs) / M.n(c, M.PRE, "D", zs))
            for k, zs in (("station", W.STATION_ZONES), ("other", W.OTHER_ZONES), ("all", W.CITY_ZONES))}
    share_st = (M.n(c, M.PRE, "P", W.STATION_ZONES) / M.n(c, M.PRE, "P"),
                M.n(c, M.PRE, "D", W.STATION_ZONES) / M.n(c, M.PRE, "D"))
    check("R1 dies to the data: drop-offs fell %.1f per cent in the station zones and %.1f in the rest (%.1f overall), "
          "and the station zones carry %.2f of pickups against %.2f of drop-offs before the move" % (
              fall["station"], fall["other"], fall["all"], share_st[0], share_st[1]),
          fall["station"] > 25 and fall["other"] < 0 and share_st[1] - share_st[0] > 0.15)
    dist = {}
    for z, name, district, km in W.ZONES:
        if z != W.AIRPORT:
            dist.setdefault(district, []).append(z)
    off = {"by the zone list's districts": solve(c, amount, groups=dist),
           "Central station's zone alone": solve(c, amount, groups={"s": [24], "o": [z for z in W.CITY_ZONES if z != 24]}),
           "the line's zones with Old Town and Market District": solve(c, amount, groups={
               "s": [2, 14, 21, 22, 24, 25], "o": [z for z in W.CITY_ZONES if z not in (2, 14, 21, 22, 24, 25)]})}
    check("off the ladder: %s" % "; ".join("%s $%.2f (%.1f per cent)" % (k, r["fee"], r["share"]) for k, r in off.items()),
          off["by the zone list's districts"]["fee"] == 6.75 and off["Central station's zone alone"]["fee"] == 6.75
          and off["the line's zones with Old Town and Market District"]["pct"] != g["pct"])
    pax = passengers(target)
    pre_m = ["2023-11", "2023-12"] + ["2024-%02d" % m for m in range(1, 11)]
    dep_pre = sum(pax[m][1] for m in pre_m)
    dep_25 = sum(pax[m][1] for m in M25)
    p_pre = M.n(c, (D("2023-11-01"), D("2024-10-31")), "P")
    q_pax = dep_25 * p_pre / dep_pre
    check("off the ladder: pickups per arriving passenger (the distractor workbook) fell %.1f per cent and restate "
          "to %.0f and $%.2f" % (100 * (1 - g["rec"] / q_pax), q_pax, M.fee_for(amount, q_pax)),
          M.fee_for(amount, q_pax) not in (6.50, 7.00, 7.25) and 23 < 100 * (1 - g["rec"] / q_pax) < 27)
    lo, hi = M.flip(g["rec"], amount, g["fee"])
    check("flip: the fee holds for any share from %.1f to %.1f per cent; the golden %.2f sits %.1f and %.1f points "
          "inside" % (100 * lo, 100 * hi, g["share"], g["share"] - 100 * lo, 100 * hi - g["share"]),
          F["flip_pct"] == {"lo": "%.1f" % (100 * lo), "hi": "%.1f" % (100 * hi)} and
          min(g["share"] - 100 * lo, 100 * hi - g["share"]) > 1.5)
    variants = [g["q"], zq, zq0] + [r["q"] for r in wr.values()]
    sur = {M.round_to(6.50 * v - amount, 10000) for v in variants}
    sho = {M.round_to(amount - 6.25 * v, 10000) for v in variants}
    check("bins: the surplus at $6.50 (%.0f) and the shortfall at $6.25 (%.0f) round to $10,000 under every window and "
          "grouping (%d variants)" % (6.50 * g["q"] - amount, amount - 6.25 * g["q"], len(variants)),
          sur == {10000} and sho == {10000} and F["surplus_at_fee"] == "10,000" and
          F["shortfall_one_step_below"] == "10,000")
    check("bins: the restated figure sits %.0f from its nearest-thousand edge and the share %.2f points from its "
          "whole-per-cent edge under every variant" % (min(abs(v - (M.round_to(v, 1000) + s * 500)) for v in variants
                                                          for s in (-1, 1)),
                                                      min(abs(100 * (1 - g["rec"] / v) - (11 + s * 0.5))
                                                          for v in variants for s in (-1, 1))),
          all(M.round_to(v, 1000) == 84000 and M.round_to(100 * (1 - g["rec"] / v), 1) == 11 for v in variants))
    sett = M.settings(amount, g["q"])
    check("the 25-cent steps (figures.json only): $5.50 to $6.25 do not recover the amount, $6.50 to $7.25 do",
          [s["recovers"] for s in sett] == [False] * 4 + [True] * 4 and
          [s["recovers"] for s in F["settings"]] == ["no"] * 4 + ["yes"] * 4)

    # ------------------------------------------------------------------ the asks: devices, stops, fairness
    gold = M.ledger(B.ledger_pickups(target)).set_index(["month", "brokerage"])
    tp = tr[tr["dir"] == "P"]
    on = pd.to_datetime(tp["meter_on"])
    keep = ((on >= "2025-01-01") & (on < "2026-01-01")) | ((on >= "2026-01-01") &
                                                           (pd.to_datetime(tp["received"]) < "2026-09-01"))
    truth_l = ledger_of(tp["meter_on"][keep], tp["brokerage"][keep])
    check("asks fair: the golden ledger equals the clean truth's (every pickup billed to its brokerage when it began), "
          "%d brokerage-months" % len(gold), gold.equals(truth_l) and len(gold) == 120)
    p25 = t25[t25["pu_zone"] == W.AIRPORT]
    sp = B.spells(target)
    lab = ledger_of(p25["meter_on"], p25["brokerage_id"])
    first = sp.sort_values("from").groupby("vehicle_id").first()["brokerage_id"]
    at0 = sp[(sp["from"] <= "2025-01-01") & (sp["to"] >= "2025-01-01")].set_index("vehicle_id")["brokerage_id"]
    ros = ledger_of(p25["meter_on"], p25["vehicle_id"].map(at0).fillna(p25["vehicle_id"].map(first)))
    cur = sp.sort_values("from").groupby("vehicle_id").last()["brokerage_id"]
    curl = ledger_of(p25["meter_on"], p25["vehicle_id"].map(cur))
    a1, a2, a3 = misses(lab, gold, M25), misses(ros, gold, M25), misses(curl, gold, M25)
    check("Ask A stops: the annual file's brokerage (the cracker) misses %d of %d brokerage-months, the roster at "
          "1 January 2025 %d, the register's current brokerage %d" % (a1[0], a1[1], a2[0], a3[0]),
          a1 == (72, 72) and a2[0] >= 20 and a3 == (72, 72))
    lab_on = (p25["brokerage_id"].astype(int).to_numpy() != B.brokerage_at_start(p25, sp)).sum()
    check("Ask A device (D2): the 2025 file labels %d pickups with a brokerage other than the one the taxi belonged to "
          "when the trip began, every brokerage gaining and losing licences at the 2026 renewal" % lab_on,
          lab_on > 2000)
    fee_col = pd.DataFrame({"month": p25["meter_on"].dt.strftime("%Y-%m"), "brokerage": B.brokerage_at_start(p25, sp),
                            "owed": p25["airport_fee"]}).groupby(["month", "brokerage"])["owed"].sum()
    a4 = misses(fee_col.to_frame("owed"), gold, M25, col="owed")
    check("Ask A hazard: fees owed summed from the meter's fee column miss January's %d brokerages (meters charged "
          "$5.50 until reprogrammed; the by-law bills the fee in force)" % a4[0], a4[0] == 6)
    cols_x = [k for k in d26nv.columns if k not in ("record_id", "drop")]
    crack = d26nv[~d26nv.duplicated(subset=cols_x)]
    pk = crack[crack["pu_zone"] == W.AIRPORT]
    b1 = misses(ledger_of(pk["meter_on"], pk["brokerage_id"]), gold, M26)
    u = B.drop_copies(B.drop_voids(rawu))
    pu = u[u["pu_zone"] == W.AIRPORT]
    b2 = misses(ledger_of(pu["meter_on"], pu["brokerage_id"]), gold, M26)
    vk = B.drop_copies(d26raw[d26raw["total"] >= 0])
    pv = vk[vk["pu_zone"] == W.AIRPORT]
    b3 = misses(ledger_of(pv["meter_on"], pv["brokerage_id"]), gold, M26)
    pdm = d26[d26["pu_zone"] == W.AIRPORT]
    dml = pd.DataFrame({"month": ["2026-" + k[-2:] for k in pdm["drop"]], "brokerage": pdm["brokerage_id"].astype(int)})
    b4 = misses(dml.groupby(["month", "brokerage"]).size(), gold, M26)
    memo_rows = (d26nv["meter_on"].dt.second == 0) & (d26nv["meter_off"].dt.second == 0) & (d26nv["sf_flag"] == "Y") & \
                (d26nv["vendor_id"] == 2)
    oc = d26nv[~memo_rows.to_numpy()]
    po = oc[oc["pu_zone"] == W.AIRPORT]
    b5 = misses(ledger_of(po["meter_on"], po["brokerage_id"]), gold, M26)
    exact = len(d26nv) - len(crack)
    check("Ask B stops: exact duplicates dropped with UTC converted and void pairs out (the cracker, %d exact copies "
          "gone) misses %d of %d, UTC months %d, voided originals kept %d, file months %d, every minute-stamped memory "
          "record dropped %d" % (exact, b1[0], b1[1], b2[0], b3[0], b4[0], b5[0]),
          b1 == (48, 48) and b2[0] >= 40 and b3[0] >= 40 and b4[0] >= 30 and b5 == (48, 48) and 20 < exact < 200)
    gone = set(d26nv["record_id"]) - set(d26["record_id"])
    resent = set(x["recs"][x["recs"]["kind"] == "resend"]["rid"])
    resent_in = resent & set(d26raw["record_id"])
    check("Ask B device (D1): the golden key (licence, start and stop minute, zones, fare, total) removes exactly the "
          "%d re-sent copies the files hold and nothing else" % len(resent_in), gone == resent_in and len(gone) > 1000)
    cp_ = d26nv[d26nv["record_id"].isin(gone) & (d26nv["pu_zone"] == W.AIRPORT)]
    per = cp_.groupby([cp_["meter_on"].dt.strftime("%Y-%m"), "brokerage_id"]).size()
    check("Ask B device (D1): every brokerage-month of 2026 holds re-sent pickups (at least %d), and the minute-stamped "
          "memory records hold %d genuine trips beside the %d copies" % (per.min(), int(memo_rows.sum()) - len(gone),
                                                                         len(gone)),
          len(per) == 48 and per.min() >= 2 and int(memo_rows.sum()) - len(gone) > 1000)
    p26 = d26[d26["pu_zone"] == W.AIRPORT]
    fee26 = pd.DataFrame({"month": p26["meter_on"].dt.strftime("%Y-%m"), "brokerage": p26["brokerage_id"].astype(int),
                          "owed": p26["airport_fee"]}).groupby(["month", "brokerage"])["owed"].sum()
    b6 = misses(fee26.to_frame("owed"), gold, M26, col="owed")
    names = {k: pq.ParquetFile(B.path(target, k)).schema_arrow.names for k in W.DROP_KEYS}
    renamed = [k for k, v in names.items() if "Airport_Fee" in v]
    check("Ask B hazards: the fee column summed misses January's %d brokerages; the column is Airport_Fee in the %d "
          "files from June" % (b6[0], len(renamed)), b6[0] == 6 and renamed == ["drop06", "drop07", "drop08"])

    # ------------------------------------------------------------------ separation: the call against the asks
    dup25 = B.drop_copies(t25)
    p24 = t24raw[t24raw["total"] > 0]
    check("separation: no re-sent copy and no void in the 2025 annual file, no re-sent copy in 2024's, and the call's "
          "windows read no 2026 record", len(dup25) == len(t25) and (t25["total"] < 0).sum() == 0 and
          len(B.drop_copies(p24)) == len(p24) and pd.Timestamp(M.BASE[1]) < d26["meter_on"].min())
    check("separation: the call's counts come only from meter start and the two zones; the asks' devices sit in the "
          "brokerage column and the 2026 files", same(solve(noc, amount)) and gold.index.get_level_values(0).min() ==
          "2025-01")

    # ------------------------------------------------------------------ the goldens against figures.json
    led = list(csv.DictReader(open(os.path.join(golden, W.GOLDEN["ledger"]), encoding="utf-8")))
    gl = [(r.month, int(r.brokerage), int(r.pickups), "%.2f" % r.owed) for r in gold.reset_index().itertuples()]
    check("golden ledger: %d rows, each equal to the recomputed ledger and to figures.json" % len(led),
          [(r["month"], int(r["brokerage_id"]), int(r["fee_liable_pickups"]), r["fees_owed"]) for r in led] == gl and
          [(r["month"], r["brokerage_id"], r["pickups_raw"], r["owed"].replace(",", "")) for r in F["ledger"]] == gl
          and all(r["brokerage_name"] == W.BROKERAGES[int(r["brokerage_id"])] for r in led))
    memo = docx_text(os.path.join(golden, W.GOLDEN["memo"]))
    need = ["Set the 2027 taxi pickup fee at $6.50", F["components"]["recorded_2025"], F["components"]["restated_2025"],
            "%s per cent" % F["components"]["share_cost_pct"], "$" + F["components"]["amount_2027"], "$10,000",
            "%s to %s per cent" % (F["flip_pct"]["lo"], F["flip_pct"]["hi"])]
    check("golden memo: states the fee, the recorded and restated pickups, the share, the amount, the surplus and "
          "shortfall and the flip, and tabulates no fee steps (the prompt asks for none)", all(n in memo for n in need)
          and not re.search(r"\$\d\.\d\d \| (Yes|No)", memo) and "|" not in memo,
          [n for n in need if n not in memo])
    with zipfile.ZipFile(os.path.join(golden, W.GOLDEN["memo"])) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
        pngb = z.read(media[0]) if media else b""
    check("golden memo: one embedded chart, written by code, no library banner", len(media) == 1 and
          pngb[:4] == b"\x89PNG" and b"Matplotlib" not in pngb and b"matplotlib" not in pngb)

    # ------------------------------------------------------------------ distractors, input gates
    tmp = copy_without(target, W.DISTRACTORS)
    try:
        fd = GO.compute(tmp)
        same_d = fd["fee"] == fx["fee"] and fd["restated"] == fx["restated"] and fd["ledger"].equals(fx["ledger"])
    finally:
        shutil.rmtree(tmp)
    check("distractors: with %s deleted the call, every component and every ledger row are unchanged" % " and ".join(
        meta["distractor_files"]), same_d and sorted(meta["distractor_files"]) == sorted(W.FILES[k] for k in W.DISTRACTORS))
    fmts = {f.rsplit(".", 1)[1] for f in files}
    rows = {f: (pq.ParquetFile(os.path.join(target, f)).metadata.num_rows if f.endswith(".parquet") else
                sum(1 for _ in open(os.path.join(target, f), encoding="utf-8")) - 1 if f.endswith(".csv") else 0)
            for f in files}
    check("input gates: %d files, %d formats (%s), the largest %s with %d rows, %d distractors named only in "
          "metadata.json" % (len(files), len(fmts), ", ".join(sorted(fmts)), max(rows, key=rows.get), max(rows.values()),
                             len(meta["distractor_files"])),
          len(files) >= 10 and len(fmts) >= 3 and max(rows.values()) >= 25000 and len(meta["distractor_files"]) >= 2
          and not any("distractor" in f.lower() for f in files) and all(f in files for f in meta["distractor_files"]))
    check("metadata: the deliverables are the two files the prompt names", meta["deliverables"] ==
          [W.GOLDEN["memo"], W.GOLDEN["ledger"]] and all(f in prompt for f in meta["deliverables"]))

    # ------------------------------------------------------------------ house fixes, wired
    dfall = {z: 100 * (1 - M.n(c, M.BASE, "D", [z]) / M.n(c, M.PRE, "D", [z])) for z in W.CITY_ZONES}
    check("H3 step 3: the four station zones are the only zones whose drop-offs fell from the twelve months before the "
          "move to 2025 (%s), each by about a third; no other zone's fell (the largest change %.1f)" % (", ".join(
              "%d %.1f" % (z, dfall[z]) for z in W.STATION_ZONES), max(dfall[z] for z in W.OTHER_ZONES)),
          all(28 < dfall[z] < 36 for z in W.STATION_ZONES) and all(dfall[z] <= 0 for z in W.OTHER_ZONES))
    mon = {m: p for m, p, _ in fx["monthly"]}
    lt = gold.reset_index().groupby("month")["pickups"].sum()
    check("H18 the ledger is the memo's series at billing grain: its monthly totals equal the chart's pickups in every "
          "month from January 2025 to August 2026, and 2025's sum is the recorded %d" % g["rec"],
          all(mon[m] == int(lt[m]) for m in M25 + M26) and int(lt[M25].sum()) == g["rec"])
    idx = [r["file_name"] for r in csv.DictReader(open(B.path(target, "index"), encoding="utf-8"))]
    check("H9 the folder list names every shipped file but itself, once", sorted(idx) == sorted(
        f for f in files if f != W.FILES["index"]) and len(idx) == len(set(idx)))
    tree = [os.path.relpath(os.path.join(r, f), TASK) for r, _, fs in os.walk(TASK) for f in fs]
    check("H11 one submission.md and one prompt.txt under the task, no superseded copy or snapshot",
          sum(os.path.basename(t) == "submission.md" for t in tree) == 1 and
          sum(os.path.basename(t) == "prompt.txt" for t in tree) == 1 and
          not any(re.search(r"_v\d|\.bak$|snapshot|DESIGN_V|prompt\.md", t) for t in tree))

    # ------------------------------------------------------------------ realism sweep
    big = {f: n for f, n in rows.items() if n >= 1000}
    check("realism: no large file with a round thousand of rows and no round headline (%d recorded, $%.2f assigned)" % (
        g["rec"], amount), all(n % 1000 for n in big.values()) and g["rec"] % 100 and amount % 1000, big)
    blobs = {}
    for root in (target, golden):
        for f in os.listdir(root):
            p = os.path.join(root, f)
            data = open(p, "rb").read()
            if f.endswith((".xlsx", ".docx")):
                with zipfile.ZipFile(p) as z:
                    data = b"".join(z.read(n) for n in z.namelist())
            blobs[f] = data
    tells = {f: [t for t in (b"openpyxl", b"ReportLab", b"reportlab", b"python-docx", b"matplotlib", b"Matplotlib",
                             b"pandas") if t in d] for f, d in blobs.items()}
    check("realism: containers scrubbed (no writing library named in any file)", not any(tells.values()),
          {f: v for f, v in tells.items() if v})
    texts = dict(docs)
    texts.update({f: open(os.path.join(target, f), encoding="utf-8").read() for f in files if f.endswith(".csv")})
    texts.update({"memo": memo, "ledger": open(os.path.join(golden, W.GOLDEN["ledger"]), encoding="utf-8").read()})
    for fn in ("prompt.txt", "submission.md", "metadata.json", "DESIGN_NOTE.md"):
        texts[fn] = open(os.path.join(TASK, fn), encoding="utf-8").read()
    check("realism: no em or en dash in any shipped text, golden, prompt, submission, metadata or design note",
          not any(DASH.search(t) for t in texts.values()), [f for f, t in texts.items() if DASH.search(t)])
    people = set(re.findall(r"\b(?:Olivia|Joseph|Johnny|Jo|Ms\.|Mr\.) [A-Z][a-z]+", " ".join(docs.values()) + memo))
    check("realism: every person named comes from guard.py names (%s)" % sorted(people),
          people <= PERSONAS and set(W.PEOPLE.values()) == PERSONAS and "Joseph Heath" in note and
          "Johnny Houston" in docs[W.FILES["dictionary"]] and "Olivia White" in memo)
    import docs as DC
    asof = {W.FILES[k]: d for k, src, d, desc in DC.INDEX}
    mt = {f: dt.datetime.fromtimestamp(os.path.getmtime(os.path.join(target, f))).date().isoformat() for f in files}
    check("realism: every file's saved date is its index date (the folder list's own date otherwise), none after %s" %
          W.AS_OF, all(mt[f] == asof.get(f, mt[f]) and mt[f] <= W.AS_OF for f in files),
          {f: (mt[f], asof.get(f)) for f in files if mt[f] != asof.get(f, mt[f])})
    pdfd = {f: re.search(rb"/CreationDate \(D:(\d{8})", open(os.path.join(target, f), "rb").read()) for f in files
            if f.endswith(".pdf")}
    check("realism: every PDF's creation date is on or before its index date",
          all(m and m.group(1).decode() <= asof[f].replace("-", "") for f, m in pdfd.items()),
          {f: m.group(1) if m else None for f, m in pdfd.items()})
    last = {k: d26raw[d26raw["drop"] == k]["meter_on"].max().date().isoformat() for k in W.DROP_KEYS}
    check("realism: each monthly file holds no trip begun after its month and is published after it",
          all(last[k] < "2026-%02d-01" % (int(k[-2:]) + 1) if int(k[-2:]) < 12 else True for k in W.DROP_KEYS) and
          all(asof[W.FILES[k]] > "2026-%02d-28" % int(k[-2:]) for k in W.DROP_KEYS))
    led_counts = [r[2] for r in gl]
    check("realism: brokerage-month pickups vary (%d distinct values in %d rows)" % (len(set(led_counts)), len(gl)),
          len(set(led_counts)) > 100)
    print("%d checks failed" % len(bad) if bad else "all checks passed")
    return not bad


if __name__ == "__main__":
    ok = run(os.path.join(TASK, "target"), os.path.join(TASK, "golden"))
    sys.exit(0 if ok else 1)
