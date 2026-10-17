"""The independent verifier: recompute the call, the restated base year, the rank's share, every fee setting, the surplus
and shortfall and every brokerage-month of the ledger from target/ alone, with its own reading of the files (pyarrow,
pdftotext, openpyxl, no generator module), and compare them with the goldens, figures.json and submission.md.

    python3 task467/generator/verify.py
"""
import csv
import datetime as dt
import json
import math
import os
import re
import subprocess
import sys
import zipfile

import numpy as np
import pandas as pd
import pyarrow.parquet as pq

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
LOCAL = "America/Toronto"


def text(path):
    out = subprocess.run(["pdftotext", "-raw", path, "-"], capture_output=True, text=True, check=True).stdout
    return re.sub(r"\s+", " ", out)


def find(target, pattern):
    hits = sorted(f for f in os.listdir(target) if re.search(pattern, f))
    assert hits, pattern
    return [os.path.join(target, h) for h in hits]


# ---------------------------------------------------------------- the trip records
def read_trips(path, names):
    t = pq.read_table(path).to_pandas()
    t = t.rename(columns={k: names.get(k, k) for k in t.columns})
    t = t.rename(columns={k: "airport_fee" for k in t.columns if k.lower() == "airport_fee"})
    for k in ("meter_on", "meter_off"):
        if t[k].dt.tz is not None:
            t[k] = t[k].dt.tz_convert(LOCAL).dt.tz_localize(None)
    return t


def net_voids(t):
    """Each reversal (amounts negated) cancels one record with the same licence, times, zones and total."""
    key = ["vehicle_id", "meter_on", "meter_off", "pu_zone", "do_zone"]
    t = t.assign(_abs=t["total"].abs().round(2), _sign=np.sign(t["total"]))
    t = t.sort_values("record_id", kind="stable")
    t["_k"] = t.groupby(key + ["_abs", "_sign"]).cumcount()
    pos, neg = t[t["_sign"] > 0], t[t["_sign"] < 0]
    hit = pos.merge(neg[key + ["_abs", "_k"]], on=key + ["_abs", "_k"], how="left", indicator=True)
    keep = hit[hit["_merge"] == "left_only"]
    assert len(t[t["_sign"] == 0]) == 0
    return keep.drop(columns=["_abs", "_sign", "_k", "_merge"])


def drop_resent(t):
    """A record whose licence, start and stop minute, zones, fare and total match an earlier record number is a copy."""
    t = t.sort_values("record_id", kind="stable")
    k = pd.DataFrame({"v": t["vehicle_id"].to_numpy(), "a": t["meter_on"].dt.floor("min").to_numpy(),
                      "b": t["meter_off"].dt.floor("min").to_numpy(), "p": t["pu_zone"].to_numpy(),
                      "d": t["do_zone"].to_numpy(), "f": t["fare"].round(2).to_numpy(),
                      "x": t["total"].round(2).to_numpy()})
    return t[~k.duplicated().to_numpy()]


def daily(t):
    pick = t["pu_zone"] == 1
    return pd.DataFrame({"date": t["meter_on"].dt.normalize(), "dir": np.where(pick, "P", "D"),
                         "zone": np.where(pick, t["do_zone"], t["pu_zone"]), "n": 1})


def window(c, a, b, d, zones=None):
    m = (c["date"] >= pd.Timestamp(a)) & (c["date"] <= pd.Timestamp(b)) & (c["dir"] == d)
    if zones is not None:
        m &= c["zone"].isin(zones)
    return int(c.loc[m, "n"].sum())


def lowest_step(amount, pickups):
    k = 1
    while k * 25 * pickups < amount * 100:
        k += 1
    return k * 25 / 100.0


def main(target, golden):
    names_v2 = {"pickup_datetime": "meter_on", "dropoff_datetime": "meter_off", "pickup_zone": "pu_zone",
                "dropoff_zone": "do_zone", "passenger_count": "passengers", "trip_distance_km": "distance_km",
                "fare_amount": "fare", "tip_amount": "tip", "total_amount": "total", "store_and_fwd_flag": "sf_flag"}
    dic = text(find(target, r"data_dictionary\.pdf$")[0])
    for a, b in names_v2.items():
        assert "%s %s" % (a, b) in dic, (a, b)
    y24 = net_voids(read_trips(find(target, r"airport_2024\.parquet$")[0], names_v2))
    y25 = net_voids(read_trips(find(target, r"airport_2025\.parquet$")[0], names_v2))
    drops = pd.concat([read_trips(p, names_v2) for p in find(target, r"airport_2026-\d\d\.parquet$")], ignore_index=True)
    y26 = drop_resent(net_voids(drops))
    arch = pd.read_csv(find(target, r"archive_2023\.csv$")[0])
    arch = pd.DataFrame({"date": pd.to_datetime(arch["trip_date"]), "dir": np.where(arch["pu_zone"] == 1, "P", "D"),
                         "zone": np.where(arch["pu_zone"] == 1, arch["do_zone"], arch["pu_zone"]), "n": arch["trips"]})
    c = pd.concat([arch, daily(y24), daily(y25), daily(y26)], ignore_index=True)

    # the line's station zones: the bulletin's stations against the zone list's names
    bul = text(find(target, r"airport_line_service_bulletin\.pdf$")[0])
    m = re.search(r"between (\w+) and Airport stations, calling at ([\w ,]+?)\.", bul)
    stations = [m.group(1)] + [s.strip() for s in re.split(r",| and ", m.group(2)) if s.strip()]
    zl = pd.read_csv(find(target, r"taxi_zones\.csv$")[0])
    st_z = sorted(int(r.zone_id) for r in zl.itertuples() if r.zone_name in stations or
                  any(r.zone_name.endswith(" " + s) for s in stations))
    city = [int(z) for z in zl["zone_id"] if z != 1]
    rest = [z for z in city if z not in st_z]
    assert st_z == [2, 14, 21, 24], (stations, st_z)

    pre, base = ("2023-11-04", "2024-11-03"), ("2025-01-01", "2025-12-31")
    rec = window(c, *base, "P")
    rec_file = int((y25["pu_zone"] == 1).sum())
    q = sum(window(c, *base, "D", zs) * window(c, *pre, "P", zs) / window(c, *pre, "D", zs) for zs in (st_z, rest))
    q_kept = sum(window(c, *base, "P", zs) / ((window(c, *base, "P", zs) / window(c, *pre, "P", zs)) /
                                              (window(c, *base, "D", zs) / window(c, *pre, "D", zs)))
                 for zs in (st_z, rest))
    wb = __import__("openpyxl").load_workbook(find(target, r"ground_transportation_budget\.xlsx$")[0], data_only=True)
    amount = [r[2] for r in wb.active.iter_rows(values_only=True) if r[1] == "Amount assigned to the taxi pickup fee"][0]
    fee = lowest_step(amount, q)
    share = 1 - rec / q
    own = lowest_step(amount, window(c, *pre, "P"))
    agg = lowest_step(amount, rec / ((rec / window(c, *pre, "P")) / (window(c, *base, "D") / window(c, *pre, "D"))))
    surplus, shortfall = fee * q - amount, amount - (fee - 0.25) * q

    def thou(v, unit):
        return int(math.floor(v / unit + 0.5)) * unit

    print("verify: 2025 fee-liable pickups %d (the file's own count %d)" % (rec, rec_file))
    print("verify: restated %.1f (%.1f by the kept share), %d to the thousand; the rank's share %.2f per cent" % (
        q, q_kept, thou(q, 1000), 100 * share))
    print("verify: amount $%.2f; fee $%.2f; the pickups' own fall would give $%.2f, all pickups against all drop-offs "
          "$%.2f" % (amount, fee, own, agg))
    print("verify: surplus at $%.2f %.0f, shortfall at $%.2f %.0f" % (fee, surplus, fee - 0.25, shortfall))
    F = json.load(open(os.path.join(HERE, "figures.json"), encoding="utf-8"))
    assert rec == rec_file and F["components"]["recorded_2025"] == "{:,}".format(rec) and abs(q - q_kept) < 1e-6
    assert F["components"]["restated_2025"] == "{:,}".format(thou(q, 1000)) and thou(q, 1000) == 84000
    assert fee == 6.50 and own == 5.75 and agg == 7.00 and thou(100 * share, 1) == 11
    qz = sum(window(c, *base, "D", [z]) * window(c, *pre, "P", [z]) / window(c, *pre, "D", [z]) for z in city)
    print("verify: every zone its own group restates to %.1f, %d to the thousand, $%.2f" % (
        qz, thou(qz, 1000), lowest_step(amount, qz)))
    assert thou(qz, 1000) == 84000 and lowest_step(amount, qz) == 6.50
    assert thou(surplus, 10000) == 10000 and thou(shortfall, 10000) == 10000

    # the ledger: 2025 by the register's affiliation on the trip's start date, 2026 by the monthly files' brokerage
    law = text(find(target, r"fees_bylaw_gt3\.pdf$")[0])
    sched = {int(y): float(v) for y, v in re.findall(r"\b(20\d\d) \$(\d+\.\d\d)\b", law)}
    aff = pd.read_csv(find(target, r"vehicle_affiliations\.csv$")[0], dtype=str, keep_default_na=False)
    aff["a"] = pd.to_datetime(aff["effective_from"])
    aff["z"] = pd.to_datetime(aff["effective_to"].replace("", "2100-01-01"))
    p25 = y25[y25["pu_zone"] == 1][["vehicle_id", "meter_on"]].copy()
    p25["day"] = p25["meter_on"].dt.normalize()
    p25 = p25.sort_values("day").reset_index(drop=True)
    asof = pd.merge_asof(p25, aff.sort_values("a")[["vehicle_id", "a", "z", "brokerage_id"]], left_on="day",
                         right_on="a", by="vehicle_id", direction="backward")
    assert (asof["day"] <= asof["z"]).all() and asof["brokerage_id"].notna().all()
    p26 = y26[y26["pu_zone"] == 1]
    rows = pd.concat([pd.DataFrame({"month": asof["meter_on"].dt.strftime("%Y-%m"),
                                    "b": asof["brokerage_id"].astype(int)}),
                      pd.DataFrame({"month": p26["meter_on"].dt.strftime("%Y-%m"), "b": p26["brokerage_id"].astype(int)})])
    led = rows.groupby(["month", "b"]).size().reset_index(name="n")
    led["owed"] = ["%.2f" % round(n * sched[int(mo[:4])], 2) for mo, n in zip(led["month"], led["n"])]
    mine = [(r.month, int(r.b), int(r.n), r.owed) for r in led.itertuples()]
    gfile = [f for f in os.listdir(golden) if f.endswith(".csv")][0]
    theirs = [(r["month"], int(r["brokerage_id"]), int(r["fee_liable_pickups"]), r["fees_owed"])
              for r in csv.DictReader(open(os.path.join(golden, gfile), encoding="utf-8"))]
    assert mine == theirs, [(a, b) for a, b in zip(mine, theirs) if a != b][:5]
    print("verify: the ledger's %d brokerage-months recompute (2025 %d pickups, 2026 to August %d)" % (
        len(mine), sum(r[2] for r in mine if r[0] < "2026"), sum(r[2] for r in mine if r[0] >= "2026")))

    # the memo, figures.json and the write-up
    from docx import Document
    memo_path = [os.path.join(golden, f) for f in os.listdir(golden) if f.endswith(".docx")][0]
    d = Document(memo_path)
    body = "\n".join([p.text for p in d.paragraphs] + [" | ".join(x.text for x in r.cells) for t in d.tables
                                                        for r in t.rows])
    for s in ("$%.2f" % fee, "{:,}".format(rec), "{:,}".format(thou(q, 1000)), "%d per cent" % thou(100 * share, 1),
              "${:,}".format(int(amount)), "$10,000"):
        assert s in body, s
    assert not d.tables and "%.1f to %.1f per cent" % (100 * (1 - rec * fee / amount),
                                                        100 * (1 - rec * (fee - 0.25) / amount)) in body
    with zipfile.ZipFile(memo_path) as z:
        assert sum(n.startswith("word/media/") for n in z.namelist()) == 1
    assert F["call"]["fee_2027"] == "%.2f" % fee and F["components"]["recorded_2025"] == "{:,}".format(rec)
    assert F["components"]["restated_2025"] == "{:,}".format(thou(q, 1000))
    assert F["components"]["share_cost_pct"] == "%d" % thou(100 * share, 1)
    sub = open(os.path.join(TASK, "submission.md"), encoding="utf-8").read()
    assert "**Set the 2027 taxi pickup fee at $%.2f.**" % fee in sub
    for mo, b, n, owed in mine:
        assert re.search(r"- %s, [A-Za-z ]+: %s pickups, \$%s\n" % (mo, re.escape("{:,}".format(n)),
                                                                  re.escape("{:,.2f}".format(float(owed)))), sub), (mo, b)
    print("verify: the memo, figures.json and submission.md carry the same fee, base year, share, flip and ledger")
    return True


if __name__ == "__main__":
    main(os.path.join(TASK, "target"), os.path.join(TASK, "golden"))
