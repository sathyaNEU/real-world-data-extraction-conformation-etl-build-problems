"""Read the shipped pack as a solver would: the Office's annual files and 2026 drops conformed to one table of trips in
local time (voided trips with their reversals out, re-sent copies out), the 2023 archive's daily counts, and each
licence's brokerage on each day from the affiliation register."""
import csv
import datetime as dt
import os

import numpy as np
import pandas as pd
import pyarrow.parquet as pq

import world as W

V2_TO_V3 = {"pickup_datetime": "meter_on", "dropoff_datetime": "meter_off", "pickup_zone": "pu_zone",
            "dropoff_zone": "do_zone", "passenger_count": "passengers", "trip_distance_km": "distance_km",
            "fare_amount": "fare", "tip_amount": "tip", "total_amount": "total", "store_and_fwd_flag": "sf_flag"}


def path(target, key):
    return os.path.join(target, W.FILES[key])


def _read(p):
    t = pq.read_table(p).to_pandas()
    t = t.rename(columns=V2_TO_V3)
    t = t.rename(columns={c: "airport_fee" for c in t.columns if c.lower() == "airport_fee"})
    for c in ("meter_on", "meter_off"):
        s = t[c]
        if getattr(s.dt, "tz", None) is not None:
            t[c] = s.dt.tz_convert(W.TZ).dt.tz_localize(None)
    return t


def drop_voids(t):
    """A reversal carries its trip's vehicle, times and zones with every amount negated: both rows go."""
    rev = t[t["total"] < 0]
    if not len(rev):
        return t
    key = ["vehicle_id", "meter_on", "meter_off", "pu_zone", "do_zone"]
    orig = t[t["total"] > 0].merge(rev[key + ["total"]].assign(total=-rev["total"]), on=key + ["total"],
                                   how="inner")["record_id"]
    out = t[(t["total"] >= 0) & ~t["record_id"].isin(orig)]
    assert len(t) - len(out) == 2 * len(rev), (len(t), len(out), len(rev))
    return out


def drop_copies(t):
    """A record matching an earlier-received one on the licence, the meter's start and stop minute, the zones and the
    amounts is a copy of it."""
    t = t.sort_values("record_id", kind="stable")
    key = pd.DataFrame({"v": t["vehicle_id"], "on": t["meter_on"].dt.floor("min"), "off": t["meter_off"].dt.floor("min"),
                        "pu": t["pu_zone"], "do": t["do_zone"], "fare": t["fare"], "total": t["total"]})
    return t[~key.duplicated(keep="first").to_numpy()]


def annual(target, year):
    return drop_voids(_read(path(target, "trips%d" % year)))


def drops(target, copies=False):
    t = pd.concat([_read(path(target, k)).assign(drop=k) for k in W.DROP_KEYS], ignore_index=True)
    t = drop_voids(t)
    return t if copies else drop_copies(t)


def as_trips(t):
    """dir and city zone from the zones: a pickup begins in zone 1."""
    pick = t["pu_zone"].to_numpy() == W.AIRPORT
    return t.assign(dir=np.where(pick, "P", "D"), zone=np.where(pick, t["do_zone"], t["pu_zone"]))


def archive(target):
    rows = list(csv.DictReader(open(path(target, "archive"), encoding="utf-8")))
    pu = np.array([int(r["pu_zone"]) for r in rows])
    do = np.array([int(r["do_zone"]) for r in rows])
    return pd.DataFrame({"date": np.array([r["trip_date"] for r in rows], dtype="datetime64[D]"),
                         "dir": np.where(pu == W.AIRPORT, "P", "D"), "zone": np.where(pu == W.AIRPORT, do, pu),
                         "n": [int(r["trips"]) for r in rows]})


def counts(target):
    """The conformed counts table, 2023 to August 2026."""
    import methods as M
    parts = [archive(target)]
    for y in (2024, 2025):
        parts.append(M.counts_from_trips(as_trips(annual(target, y))))
    parts.append(M.counts_from_trips(as_trips(drops(target))))
    c = pd.concat(parts, ignore_index=True)
    return c.groupby(["date", "dir", "zone"], as_index=False)["n"].sum()


def spells(target):
    a = pd.read_csv(path(target, "affiliations"), dtype=str, keep_default_na=False)
    a["from"] = pd.to_datetime(a["effective_from"])
    a["to"] = pd.to_datetime(a["effective_to"].replace("", "2099-12-31"))
    a["brokerage_id"] = a["brokerage_id"].astype(int)
    return a


def brokerage_at_start(trips, sp):
    """Each trip's brokerage on the day its meter started, from the register's spells."""
    day = trips["meter_on"].dt.normalize()
    m = trips[["vehicle_id"]].assign(day=day.to_numpy(), i=np.arange(len(trips))).merge(
        sp[["vehicle_id", "brokerage_id", "from", "to"]], on="vehicle_id", how="left")
    m = m[(m["from"] <= m["day"]) & (m["day"] <= m["to"])]
    assert m["i"].is_unique and len(m) == len(trips)
    return m.sort_values("i")["brokerage_id"].to_numpy()


def ledger_pickups(target):
    """Fee-liable pickups, 2025 and 2026 to date, with the brokerage each was billed to."""
    sp = spells(target)
    t25 = annual(target, 2025)
    t25 = t25[t25["pu_zone"] == W.AIRPORT]
    t26 = drops(target)
    t26 = t26[t26["pu_zone"] == W.AIRPORT]
    return pd.concat([pd.DataFrame({"meter_on": t25["meter_on"], "brokerage": brokerage_at_start(t25, sp)}),
                      pd.DataFrame({"meter_on": t26["meter_on"], "brokerage": t26["brokerage_id"].astype(int)})],
                     ignore_index=True)


def budget_amount(target):
    from openpyxl import load_workbook
    ws = load_workbook(path(target, "budget"), data_only=True).active
    for row in ws.iter_rows(values_only=True):
        if row[1] and str(row[1]).startswith("Amount assigned to the taxi pickup fee"):
            return float(row[2])
    raise ValueError("no amount line")
