"""The goldens, from the shipped files only: venue_list_board_paper_2027.pdf (the board paper: the venues kept off the
platform's list, each venue's adverse move as the platform's orders would meet it, the chart, the platform's market
orders and its 2026 dealing), venue_review_2027.xlsx (the same tables), submission.md and generator/figures.json
(every figure the write-up states, as it states it)."""
from __future__ import annotations

import io
import json
import os

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import calc as C
import common as K
import postprocess as PP

XLSX = "venue_review_2027.xlsx"
PDF = "venue_list_board_paper_2027.pdf"
POINT_NAMES = ["09:30", "12:00", "14:30"]
QUARTERS = ["January to March", "April to June", "July to September", "October"]
QSHORT = ["Q1", "Q2", "Q3", "October"]
SHORT = {v: K.VENUE[v]["short"] for v in K.VENUES}
INK = "#1d2229"
INK2 = "#4b525c"
GRID = "#e3e6ea"
OFF = "#b5443a"
ON = "#3f6f93"


def fmt(x, dp=0):
    v = K.round_half_up(x, dp)
    return format(v, ",.%df" % dp) if dp > 0 else format(int(v), ",")


class Out:
    pass


def compute(target):
    P = C.load(target)
    F = C.fills(P)
    X = Out()
    X.P = P
    X.est = C.golden(P, F)
    X.excl = list(C.excluded(X.est))
    X.kept = [v for v in K.VENUES if v not in X.excl]
    X.whole = {v: C.whole_bp(X.est[v][0]) for v in K.VENUES}
    X.n_lone = {v: X.est[v][1] for v in K.VENUES}
    X.n_fills = int(len(F))
    X.raw = C.rungs(P, F)["R0_every_fill"]
    A = C.ask_a(P)
    X.a_counts, X.a_sizes = A["counts"], A["sizes"]
    X.b = C.ask_b(P)
    return X


def names(vs):
    s = [SHORT[v] for v in vs]
    return ", ".join(s[:-1]) + " and " + s[-1] if len(s) > 1 else s[0]


# ------------------------------------------------------------------------------------------------- the chart
def chart_png(X):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    for f in ("/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf",
              "/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf"):
        font_manager.fontManager.addfont(f)
    plt.rcParams["font.family"] = "Carlito"
    plt.rcParams["svg.hashsalt"] = "task271"
    order = sorted(K.VENUES, key=lambda v: -X.est[v][0])
    vals = [X.est[v][0] for v in order]
    fig = plt.figure(figsize=(8.4, 4.1), dpi=170)
    ax = fig.add_axes([0.075, 0.22, 0.9, 0.66])
    cols = [OFF if v in X.excl else ON for v in order]
    xs = np.arange(len(order))
    ax.bar(xs, vals, width=0.62, color=cols, zorder=3)
    ax.axhline(K.LINE, color=INK, lw=1.2, ls=(0, (5, 3)), zorder=4)
    ax.text(len(order) - 0.55, K.LINE + 0.06, "Policy line: 1.5 bp", ha="right", va="bottom", fontsize=9.5, color=INK)
    for x, v, val in zip(xs, order, vals):
        ax.text(x, val + 0.06 if val >= 0 else val - 0.06, "%d bp" % X.whole[v], ha="center",
                va="bottom" if val >= 0 else "top", fontsize=9.5, color=INK)
    ax.set_xticks(xs)
    ax.set_xticklabels([SHORT[v].replace(" ", "\n", 1) for v in order], fontsize=9.2, color=INK2)
    ax.set_ylabel("Adverse move, 60 s (bp)", fontsize=10, color=INK2)
    ax.set_ylim(min(0, min(vals)) - 0.4, max(vals) + 0.6)
    ax.grid(axis="y", color=GRID, lw=0.7, zorder=1)
    ax.tick_params(axis="y", labelsize=9.5, colors=INK2, length=0)
    ax.tick_params(axis="x", length=0)
    for sp in ("top", "right", "left"):
        ax.spines[sp].set_visible(False)
    ax.spines["bottom"].set_color("#9aa1aa")
    fig.text(0.075, 0.955, "Four venues fail the test for the platform's orders", fontsize=12, color=INK,
             fontweight="bold", va="top")
    fig.text(0.075, 0.035, "Red: kept off the list. Each bar is the average adverse move over the desk's 2026 fills with "
             "no desk trade in the same share in the following minute.", fontsize=8.6, color=INK2)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=170, metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


# ------------------------------------------------------------------------------------------------- the workbook
def write_xlsx(path, X):
    wb = Workbook()
    thin = Side(style="thin", color="C9CED4")
    head_fill = PatternFill("solid", fgColor="E8ECF0")
    bold = Font(name="Calibri", bold=True)

    def header(ws, row, cols):
        for j, c in enumerate(cols, 1):
            cell = ws.cell(row=row, column=j, value=c)
            cell.font = bold
            cell.fill = head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=thin)

    ws = wb.active
    ws.title = "Venue list"
    ws["A1"] = "Platform venue list from 4 January 2027: the adverse-move test (policy 4.3)"
    ws["A1"].font = Font(name="Calibri", bold=True, size=12)
    ws["A2"] = ("Adverse move: the mid 60 seconds after a fill against the mid at the fill, basis points, positive "
                "against the order. Read on the desk's 2026 fills with no desk trade in the same share in the "
                "following minute, busted executions out; line above 1.5 bp.")
    header(ws, 4, ["Venue", "MIC", "Type", "Fills read", "Adverse move for the platform's orders (bp)", "Platform list"])
    r = 5
    for v in K.VENUES:
        mic = v if v != "THGD" else "THGD / THGM"
        ws.cell(row=r, column=1, value=K.VENUE[v]["name"])
        ws.cell(row=r, column=2, value=mic)
        ws.cell(row=r, column=3, value=K.VENUE[v]["kind"])
        c = ws.cell(row=r, column=4, value=X.n_lone[v])
        c.number_format = "#,##0"
        ws.cell(row=r, column=5, value=X.whole[v]).number_format = "0"
        ws.cell(row=r, column=6, value="Kept off" if v in X.excl else "On the list")
        if v in X.excl:
            for j in range(1, 7):
                ws.cell(row=r, column=j).font = Font(name="Calibri", bold=True, color="9C2B22")
        r += 1
    for col, w in zip("ABCDEF", (32, 13, 24, 11, 22, 13)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"

    ws2 = wb.create_sheet("Market orders 2027")
    ws2["A1"] = "The platform's 2026 dealt orders under the 2027 dealing plan: market orders at each dealing point"
    ws2["A1"].font = Font(name="Calibri", bold=True, size=12)
    ws2["A2"] = ("Client orders Gresmere dealt with trade dates 2 January to 30 October 2026, each at the point its "
                 "placing time gives under the plan's 15-minute cut-offs, whole; buys and sells in a share netted.")
    header(ws2, 4, ["Dealing point", "Mean size, % of the share's ADV"])
    for i, pn in enumerate(POINT_NAMES):
        ws2.cell(row=5 + i, column=1, value=pn)
        ws2.cell(row=5 + i, column=2, value=round(X.a_sizes[i], 3)).number_format = "0.000"
    for col, w in zip("AB", (14, 28)):
        ws2.column_dimensions[col].width = w

    ws3 = wb.create_sheet("Gresmere 2026")
    ws3["A1"] = "Net value the platform's clients bought through Gresmere in 2026, before charges and stamp duty"
    ws3["A1"].font = Font(name="Calibri", bold=True, size=12)
    header(ws3, 3, ["Month", "Bought less sold (£m)"])
    for m in range(1, 11):
        ws3.cell(row=3 + m, column=1, value="%s 2026" % K.MONTH_NAMES[m])
        ws3.cell(row=3 + m, column=2, value=round(X.b[m], 2)).number_format = "#,##0.00"
    for col, w in zip("AB", (18, 22)):
        ws3.column_dimensions[col].width = w

    ws4 = wb.create_sheet("Notes")
    notes = [
        "Sources: Hessenby Analytics' execution export (2 January to 30 October 2026), the Ashlowe and Pemberwick "
        "execution reports for busts, the OMS orders, the venue register, the platform's client orders and Gresmere's "
        "contract notes settling 2 January to 3 November 2026.",
        "Thursgate's midpoint book traded as THGD to 29 May and THGM from 1 June 2026; both codes are one venue here.",
        "Closing auction executions and fills in the last minute before the close have no 60-second mid and are not "
        "read.",
        "Market order sizes are in the share's current units: the consolidating line's old shares (before 18 May 2026) "
        "are divided by ten.",
        "Prepared by Declan Davies, execution research, for Owen Porter, 20 November 2026.",
    ]
    for i, t in enumerate(notes, 1):
        ws4.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True, vertical="top")
    ws4.column_dimensions["A"].width = 110
    wb.save(path)


# ------------------------------------------------------------------------------------------------- the board paper
def write_pdf(path, X, png):
    from reportlab import rl_config
    rl_config.invariant = 1
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    import docs as D
    st = D.styles("Carlito", "Carlito-Bold", colors.HexColor("#4a2f5c"), 10.4)
    s = []
    s.append(Paragraph("Fenmarsh Wealth Limited: Board, 9 December 2026. Paper FW-2026-52", st["sub"]))
    s.append(Paragraph(D.esc("Keep %s off the platform's venue list" % names(X.excl)), st["h1"]))
    s.append(Paragraph(D.esc("From Owen Porter, Head of Dealing, Fenmarsh Investment Management. 20 November 2026. "
                             "For decision."), st["sub"]))
    rec = ("When the central desk starts dealing for the platform on 4 January 2027, the platform's market orders "
           "should go only to %s. %s fail the order execution policy's adverse-move test for the platform's orders: "
           "the mid moves against an order filled there by %s basis points on average in the minute after the fill, "
           "against the policy's line of 1.5." % (names(X.kept), names(X.excl),
                                                  ", ".join("%d (%s)" % (X.whole[v], SHORT[v]) for v in X.excl)))
    s.append(Paragraph(D.esc(rec), st["body"]))
    s.append(Paragraph(D.esc("The evidence"), st["h2"]))
    ev = ("The platform's market orders are small and each goes whole to one venue in one execution, so nothing of "
          "the desk's follows them into the next minute. Most of the desk's 2026 fills were not like that: they were "
          "child orders of the algorithms' parent orders, and the same order's next fills, often within seconds, kept "
          "moving the mid the desk's way inside the earlier fills' minute. That flatters the venues the algorithms use "
          "at the start of an order, the midpoint books above all, whose averages over every 2026 fill sit under the "
          "line. Read on the %s fills with no desk trade in the same share in the following minute, the venues' moves "
          "are those in the table." % fmt(sum(X.n_lone.values())))
    s.append(Paragraph(D.esc(ev), st["body"]))
    rows = [("Venue", "Type", "Fills read", "Adverse move (bp)", "Platform list")]
    for v in K.VENUES:
        rows.append((K.VENUE[v]["name"] + (" (THGD, THGM)" if v == "THGD" else ""), K.VENUE[v]["kind"],
                     fmt(X.n_lone[v]), "%d" % X.whole[v], "Kept off" if v in X.excl else "On the list"))
    s.append(D.grid(rows, [5.6, 3.4, 2.0, 2.4, 2.4], st))
    s.append(Paragraph(D.esc("Busted executions out; Thursgate's two segment codes are one venue; closing auction fills "
                             "and fills in the last minute before the close have no 60-second mid. Source: Hessenby "
                             "Analytics' execution export, 2 January to 30 October 2026."), st["small"]))
    s.append(Spacer(1, 4))
    s.append(Image(io.BytesIO(png), width=16.0 * cm, height=16.0 * cm * 4.1 / 8.4))
    s.append(Paragraph(D.esc("The platform's market orders"), st["h2"]))
    flow = ("Under the 2027 dealing plan the client orders Gresmere dealt in 2026 come to market orders that are small "
            "against each share's average daily volume, largest at the 09:30 point, which carries the evening's and "
            "the early morning's orders.")
    s.append(Paragraph(D.esc(flow), st["body"]))
    rows = [("Dealing point", "Mean size, % of ADV")]
    for i, pn in enumerate(POINT_NAMES):
        rows.append((pn, "%.3f" % X.a_sizes[i]))
    s.append(D.grid(rows, [3.0, 3.6], st))
    s.append(Paragraph(D.esc("The platform's dealing through Gresmere in 2026"), st["h2"]))
    rows = [("Month", "Bought less sold (£m)")]
    for m in range(1, 11):
        rows.append(("%s 2026" % K.MONTH_NAMES[m], "%.2f" % X.b[m]))
    s.append(D.grid(rows, [4.0, 4.0], st))
    s.append(Paragraph(D.esc("Value of the shares at the dealing price, before commission, PTM levy and stamp duty, by "
                             "trade date; Gresmere's notes settling 2 January to 3 November 2026."), st["small"]))
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=1.8 * cm,
                            bottomMargin=1.8 * cm, title="Platform venue list from 4 January 2027",
                            author=K.PEOPLE["dealing"], subject="", creator="", producer="")
    foot = D._footer("Fenmarsh Wealth Limited  |  Board paper FW-2026-52")
    doc.build(s, onFirstPage=foot, onLaterPages=foot)
    PP.normalize_pdf(path, b"D:20261120164500+00'00'", b"D:20261120164500+00'00'")


# ------------------------------------------------------------------------------------------------- the write-up
def submission(X):
    s = []
    s += ["## Tags", "",
          "**Domain:** Quantitative Finance (market microstructure and execution cost: each venue's adverse move in the "
          "minute after a fill, read for single whole orders from a year of an asset manager's algorithmic fills).",
          "**Analytical objective:** Descriptive & Distribution Analysis.", ""]
    s += ["## 1. Final Recommendation", "",
          "**The board should keep %s off the platform's venue list and send the platform's orders to %s.**" % (
              names(X.excl), names(X.kept)),
          "Read on the desk's 2026 fills with no desk trade in the same share in the following minute, the four meet "
          "an average adverse move of %s basis points against the policy's 1.5. Not every 2026 fill averaged, which "
          "counts the same order's following child fills inside each fill's minute and passes %s; not the desk's "
          "2026 single orders, which never used the midpoint books; not a short horizon, which stops before the move "
          "is made." % (", ".join("%d (%s)" % (X.whole[v], SHORT[v]) for v in X.excl),
                        names([v for v in X.excl if X.raw[v][0] <= K.LINE])), ""]
    s += ["## 2. Critical Components", ""]
    for i, v in enumerate(X.excl, 1):
        s.append("%d. %s's adverse move for the platform's orders is **%d basis points**." % (i, K.VENUE[v]["name"],
                                                                                           X.whole[v]))
    s.append("")
    s += ["## 3. Step-by-Step Solution", "",
          "1. Took the executions in %s once each (one row per broker_exec_id), left out those %s and %s report as busted "
          "and the closing auction, and computed each fill's adverse move from mid_at_exec and mid_plus_60s as the "
          "adverse-move test in order_execution_policy_2025.pdf defines it: %s fills." % (
              "hessenby_tca_executions_2026-01-02_to_2026-10-30.csv", "ashlowe_execution_report_2026.csv",
              "pemberwick_execution_report_2026.xlsx", fmt(X.n_fills)),
          "2. Kept the fills after which no other desk execution in the same share, busted ones included, printed "
          "within 60 seconds, the "
          "platform's orders being sent whole and filled in one execution (wealth_dealing_plan_2027.pdf 4), with "
          "THGD and THGM as Thursgate's one book (venue_register_2026.xlsx): %s fills." % fmt(sum(X.n_lone.values())),
          "3. Averaged each venue's adverse move over those fills and kept off the list every venue above 1.5 basis "
          "points: %s." % names(X.excl),
          "4. For the market orders, took each client order with a Gresmere note dated 1 January to 30 October 2026 "
          "(cancelled notes out, the new system's second issue of an end-of-June trade the old system had noted, same "
          "order, trade date, time and shares, out, an order's parts summed) at the dealing point its placed_at, a "
          "UTC time, gives under the plan's 15-minute cut-offs and the trading calendar, netted buys and sells in "
          "each share at each point, and divided each market order by the share's ADV, the consolidating line's "
          "shares before 18 May divided by ten (corporate_actions_2026.json), the order book carrying the new line's "
          "code on every order: mean sizes %.3f, %.3f and %.3f per cent." % (X.a_sizes[0], X.a_sizes[1],
                                                                           X.a_sizes[2]),
          "5. For Gresmere's 2026 dealing, valued each effective note at shares times price_gbx, since the "
          "consideration column carries stamp duty on the old system's purchases and the client's net amount on the "
          "new system's notes, and took each trade month's purchases less its sales, January to October.",
          "6. The board keeps %s off the list." % names(X.excl), ""]
    s += ["## 4. Deliverable Answers", "", "### %s" % PDF, "",
          "1. Venues kept off the platform's list: %s." % names(X.excl),
          "2. Each venue's adverse move for the platform's orders, whole basis points:"]
    for v in K.VENUES:
        s.append("   - %s: %d, %s" % (K.VENUE[v]["name"], X.whole[v], "kept off" if v in X.excl else "on the list"))
    s.append("3. Bar chart of the eight venues' adverse moves against the 1.5 basis point line, the four kept off in "
             "red.")
    s += ["", "### %s" % XLSX, "",
          ]
    s.append("4. Mean size of the platform's 2026 dealt orders as market orders under the 2027 dealing plan, per cent "
             "of the share's average daily volume, three decimals:")
    for i, pn in enumerate(POINT_NAMES):
        s.append("   - %s: %.3f" % (pn, X.a_sizes[i]))
    s.append("5. Value bought less value sold through Gresmere by trade month, £ million, two decimals:")
    for m in range(1, 11):
        s.append("   - %s 2026: %.2f" % (K.MONTH_NAMES[m], X.b[m]))
    s.append("")
    return "\n".join(s)


def figures(X):
    return {
        "call": {"kept_off": [K.VENUE[v]["name"] for v in X.excl], "on_list": [K.VENUE[v]["name"] for v in X.kept],
                 "line_bp": K.LINE},
        "components": {K.VENUE[v]["name"]: X.whole[v] for v in X.excl},
        "steps": {"fills_read": fmt(X.n_fills), "lone_fills": fmt(sum(X.n_lone.values())), "horizon_s": K.HORIZON,
                  "cutoff_min": 15, "consolidation_ratio": 10},
        "answers": {
            "adverse_move_bp": {K.VENUE[v]["name"]: X.whole[v] for v in K.VENUES},
            "fills_read_by_venue": {K.VENUE[v]["name"]: fmt(X.n_lone[v]) for v in K.VENUES},
            "market_orders": {POINT_NAMES[i]: fmt(X.a_counts[i]) for i in range(3)},
            "mean_size_pct_adv": {POINT_NAMES[i]: "%.3f" % X.a_sizes[i] for i in range(3)},
            "net_bought_gbp_m": {K.MONTH_NAMES[m]: "%.2f" % X.b[m] for m in range(1, 11)},
        },
    }


def main(target, golden, fig_path, sub_path):
    os.makedirs(golden, exist_ok=True)
    X = compute(target)
    write_xlsx(os.path.join(golden, XLSX), X)
    PP.normalize_ooxml(os.path.join(golden, XLSX), "Declan Davies", "Owen Porter", "2026-11-17T09:12:00Z",
                       "2026-11-20T16:31:00Z", excel=True, company="Fenmarsh Investment Management")
    png = chart_png(X)
    write_pdf(os.path.join(golden, PDF), X, png)
    with open(sub_path, "w", encoding="utf-8") as fh:
        fh.write(submission(X))
    with open(fig_path, "w", encoding="utf-8") as fh:
        json.dump(figures(X), fh, indent=2, ensure_ascii=False)
        fh.write("\n")
    return X
