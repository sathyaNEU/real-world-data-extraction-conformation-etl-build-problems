"""The asks with each device's handling as a switch, so the golden, the cracker, the mirror and the checks run one
code path: protected homes (the hydrant numbers in force on each test's date, the flow available at 20 psi), market
value (the 2026 appeal decisions, the common level ratio in effect at certification), the book's water claim
payments (the adjusting firm's bordereau, reopened lines as the same claim, net of recoveries) and the book's fire
frequency by protection."""
import datetime as dt
import json
import os

import numpy as np
import pandas as pd
from scipy.spatial import cKDTree

import analysis as A
import common as K


def hydrants(target, renumber=True, at20=True):
    g = json.load(open(os.path.join(target, K.FILES["hydrants"])))
    H = pd.DataFrame([dict(f["properties"], x=f["geometry"]["coordinates"][0], y=f["geometry"]["coordinates"][1])
                      for f in g["features"]])
    F = A.read(target, "flowtests")
    by_new = {h: i for i, h in enumerate(H["HYDRANT_ID"])}
    by_old = {h: i for i, h in enumerate(H["FORMER_ID"]) if h}
    latest = {}
    for r in F.itertuples(index=False):
        d = dt.date.fromisoformat(r.test_date)
        i = (by_old.get(r.hydrant_id) if d < A.RENUMBER else by_new.get(r.hydrant_id)) if renumber else \
            by_new.get(r.hydrant_id)
        if i is None:
            continue
        if i not in latest or d > latest[i][0]:
            latest[i] = (d, int(r.static_psi), int(r.residual_psi), int(r.flow_gpm))
    q = []
    for i in range(len(H)):
        if i not in latest:
            q.append(np.nan)
        else:
            _, s, rr, f = latest[i]
            q.append(f * ((s - 20.0) / (s - rr)) ** 0.54 if at20 else float(f))
    H["avail"] = q
    return H


def protected(H, xy, filters=True):
    ok = H[H["avail"] >= 500]
    if filters:
        ok = ok[(ok["STATUS"] == "IN SERVICE") & (ok["OWNER"] == "TCWA")]
    dd, _ = cKDTree(ok[["x", "y"]].values).query(xy)
    return dd <= 1000.0


def protected_homes(target, homes, **kw):
    H = hydrants(target, kw.get("renumber", True), kw.get("at20", True))
    p = protected(H, homes[["x", "y"]].values, kw.get("filters", True))
    return {t: int(p[homes["town"].values == t].sum()) for t in K.CANDIDATES}


def market(target, homes, appeals=True, ratio="certified"):
    Ap, R = A.read(target, "appeals"), A.read(target, "clr")
    after = {r.parcel_id: int(r.assessed_total_after) for r in Ap.itertuples(index=False)
             if r.tax_year == "2026" and r.decision in ("REDUCED", "INCREASED", "SUSTAINED")} if appeals else {}
    if ratio == "certified":
        c = [float(r.common_level_ratio) for r in R.itertuples(index=False)
             if dt.date.fromisoformat(r.effective_from) <= A.CERTIFIED <= dt.date.fromisoformat(r.effective_to)][0]
    elif ratio == "roll_year":
        c = [float(r.common_level_ratio) for r in R.itertuples(index=False) if r.ratio_year == "2026"][0]
    else:
        c = 1.0
    v = np.array([after.get(p, a) for p, a in zip(homes["parcel_id"], homes["assessed"])], float) / c
    return {t: K.r(v[homes["town"].values == t].sum() / 1e6, 1) for t in K.CANDIDATES}


def payments(target, bordereau=True, reopen="combine", net=True, bmode="latest"):
    C, Pm = A.read(target, "claims"), A.read(target, "payments")
    wc = C[C["cause_code"].isin(A.WATER)].copy()
    wc["year"] = wc["date_of_loss"].str[:4].astype(int)
    base = wc[wc["reopened_from"] == ""] if reopen == "combine" else wc
    ids = set(base["claim_number"])
    Pm = Pm.copy()
    Pm["amount"] = Pm["amount"].astype(float)
    Pm["key"] = Pm["claim_number"].str.replace(r"-R\d+$", "", regex=True) if reopen == "combine" else Pm["claim_number"]
    Pm = Pm[Pm["key"].isin(ids)]
    ind = Pm[Pm["transaction_type"].isin(["INDEMNITY", "VOID"])].groupby("key")["amount"].sum()
    rec = Pm[Pm["transaction_type"] == "RECOVERY"].groupby("key")["amount"].sum()
    if bordereau:
        import openpyxl
        wb = openpyxl.load_workbook(os.path.join(target, K.FILES["bordereau"]), read_only=True)
        last, total = {}, {}
        for name in wb.sheetnames:
            for row in wb[name].iter_rows(min_row=3, values_only=True):
                if row[0]:
                    last[row[0]] = (float(row[5]), float(row[6]))
                    total[row[0]] = total.get(row[0], 0.0) + float(row[5])
        for no, (p, rr) in last.items():
            if no in ids:
                ind.loc[no] = ind.get(no, 0.0) + (p if bmode == "latest" else total[no])
                rec.loc[no] = rec.get(no, 0.0) - rr
    tot = ind.add(rec, fill_value=0.0) if net else ind
    paid = ind[ind > 0.005].index
    yr = base.set_index("claim_number")["year"]
    out = {}
    for y in K.YEARS:
        c = [k for k in paid if yr[k] == y]
        out[str(y)] = int(K.r(tot.loc[c].sum() / len(c)))
    return out


def fire(target, d, **kw):
    C = A.read(target, "claims")
    n = A.counts(C, A.FIRE)
    homes = d.drop_duplicates("parcel_id")[["parcel_id", "x", "y"]]
    H = hydrants(target, kw.get("renumber", True), kw.get("at20", True))
    prot = pd.Series(protected(H, homes[["x", "y"]].values, kw.get("filters", True)), index=homes["parcel_id"].values)
    ex = d.groupby("parcel_id")["expo"].sum()
    fc = n.groupby(level=0).sum().reindex(ex.index).fillna(0)
    out = {}
    for flag, name in ((True, "protected"), (False, "unprotected")):
        idx = prot[prot == flag].index
        out[name] = K.r(1000.0 * fc.loc[idx].sum() / ex.loc[idx].sum(), 1)
    return out


def all_asks(target, d=None, homes=None, **kw):
    if d is None:
        d, _ = A.book_table(target)
    if homes is None:
        homes = A.township_homes(target)
    pk = {k: kw[k] for k in ("renumber", "at20", "filters") if k in kw}
    return dict(protected_homes=protected_homes(target, homes, **pk),
                market_value_musd=market(target, homes, kw.get("appeals", True), kw.get("ratio", "certified")),
                water_payment_by_accident_year=payments(target, kw.get("bordereau", True), kw.get("reopen", "combine"),
                                                        kw.get("net", True), kw.get("bmode", "latest")),
                fire_per_1000=fire(target, d, **pk))
