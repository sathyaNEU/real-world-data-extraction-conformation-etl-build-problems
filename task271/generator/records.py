"""task271 records: each system's own export of the world. The cost provider's per-fill export, the two brokers'
execution reports, the OMS orders, the venue register, the securities master, the corporate actions and trading
calendar, the wealth platform's client orders and Gresmere's contract notes, and the exchange's closing auction
statistics (unused by the solution)."""
from __future__ import annotations

import datetime as dt
import json
import math

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

import common as K
import world as W
import wealth as WL

F_HESS = "hessenby_tca_executions_2026-01-02_to_2026-10-30.csv"
F_ASHL = "ashlowe_execution_report_2026.csv"
F_PEMB = "pemberwick_execution_report_2026.xlsx"
F_OMS = "oms_orders_2026.csv"
F_VENUES = "venue_register_2026.xlsx"
F_MASTER = "securities_master_2026-10-30.csv"
F_CA = "corporate_actions_2026.json"
F_CAL = "uk_trading_calendar_2025-12-01_to_2027-01-31.csv"
F_ORDERS = "wealth_platform_client_orders_2025-12-29_to_2026-10-30.csv"
F_NOTES = "gresmere_contract_notes_settling_2026-01-02_to_2026-11-03.csv"
F_AUCT = "wexcombe_closing_auction_statistics_2026.csv"

FUNDS = ["FIM UK Equity Income", "FIM UK Opportunities", "FIM UK Smaller Companies", "FIM Multi-Asset Balanced",
         "FIM Responsible UK Equity", "Fenmarsh Growth Trust plc"]


def cons_view(stock, day):
    """True when a row about the consolidated stock carries the old line's units (before 18 May)."""
    return stock["i"] == W.CONSOLIDATED_INDEX and day < W.CONSOLIDATION_DATE


def isin_of(stock, day):
    return stock["old_isin"] if cons_view(stock, day) else stock["isin"]


def sedol_of(stock, day):
    return stock["old_sedol"] if cons_view(stock, day) else stock["sedol"]


def px_view(stock, day, p):
    if p is None:
        return None
    return round(p / WL.CONS_RATIO, 4) if cons_view(stock, day) else round(p, 4)


def qty_view(stock, day, q):
    return q * WL.CONS_RATIO if cons_view(stock, day) else q


def venue_code(v, day):
    if v == "THGD" and day >= K.THURSGATE_SWITCH:
        return K.THURSGATE_NEW
    return v


def london_hms(day, sec):
    s = int(math.floor(sec))
    ms = int(round((sec - s) * 1000))
    if ms == 1000:
        s, ms = s + 1, 0
    return "%02d:%02d:%02d.%03d" % (s // 3600, (s // 60) % 60, s % 60, ms)


def prepare(X):
    """Broker, execution ids, busts and the cost provider's duplicate deliveries, all fixed once per build."""
    rg = K.rng("ids")
    stocks = X["stocks"]
    by_pid = {p.pid: p for p in X["parents"]}
    for e in X["execs"]:
        p = by_pid[e["pid"]]
        e["broker"] = "ASHL" if (p.algo == "DMA" or K.ALGO[p.algo]["broker"] == "ASHL") else "PEMB"
        e["utc"] = K.utc_of(e["day"], e["t"])
    X["execs"].sort(key=lambda e: (e["utc"], e["pid"], e["t"]))
    na = nb = 0
    for e in X["execs"]:
        if e["broker"] == "ASHL":
            na += 1
            e["exec_id"] = "ASL%09d" % (310000000 + na * 7 + int(rg.integers(0, 7)))
        else:
            nb += 1
            e["exec_id"] = "PMB-%s" % K.token("pemb", nb, n=10).upper()
    cont = [e for e in X["execs"] if e["session"] == "CONT"]
    pick = rg.random(len(cont))
    for e, r in zip(cont, pick):
        e["busted"] = bool(r < 0.0025)
    for e in X["execs"]:
        e.setdefault("busted", False)
    dup = rg.random(len(X["execs"]))
    for e, r in zip(X["execs"], dup):
        e["dup"] = bool(r < 0.003 and e["session"] == "CONT")
    for p in X["parents"]:
        p.order_id = "FIM-%s-%05d" % (p.day.strftime("%y%m%d"), p.pid % 100000)
        p.fund = FUNDS[int(K.rng("fund", p.pid).integers(0, len(FUNDS)))]
    return X


HESS_COLS = ["hessenby_ref", "client_order_id", "broker", "broker_exec_id", "isin", "side", "shares", "price_gbx",
             "last_mkt", "exec_time_utc", "session", "mid_at_exec", "mid_plus_1s", "mid_plus_5s", "mid_plus_30s",
             "mid_plus_60s", "mid_plus_300s"]


def hess_rows(X):
    stocks = X["stocks"]
    by_pid = {p.pid: p for p in X["parents"]}
    rows, n = [], 0
    for e in X["execs"]:
        s = stocks[e["stock"]]
        d = e["day"]
        m = e["marks"]
        n += 1
        r = dict(hessenby_ref="HX26%07d" % n, client_order_id=by_pid[e["pid"]].order_id, broker=e["broker"],
                 broker_exec_id=e["exec_id"], isin=isin_of(s, d), side="BUY" if e["side"] == "B" else "SELL",
                 shares=qty_view(s, d, e["qty"]), price_gbx="%.4f" % px_view(s, d, e["price"]),
                 last_mkt=venue_code(e["venue"], d), exec_time_utc=K.iso_ms(e["utc"]),
                 session="CONT" if e["session"] == "CONT" else "CLOSE_AUCTION")
        for col, h in (("mid_at_exec", 0), ("mid_plus_1s", 1), ("mid_plus_5s", 5), ("mid_plus_30s", 30),
                       ("mid_plus_60s", 60), ("mid_plus_300s", 300)):
            v = e["mid0"] if h == 0 else m.get(h)
            r[col] = "" if v is None else "%.4f" % px_view(s, d, v)
        rows.append(r)
        if e["dup"]:
            n += 1
            r2 = dict(r)
            r2["hessenby_ref"] = "HX26%07d" % n
            r2["exec_time_utc"] = e["utc"].strftime("%Y-%m-%dT%H:%M:%S.000Z")
            rows.append(r2)
    return rows


ASHL_COLS = ["exec_id", "order_id", "isin", "side", "last_qty", "last_px", "last_mkt", "transact_time", "exec_type",
             "ref_exec_id", "trading_session"]


def ashl_rows(X):
    stocks = X["stocks"]
    by_pid = {p.pid: p for p in X["parents"]}
    rows = []
    busts = []
    for e in X["execs"]:
        if e["broker"] != "ASHL":
            continue
        s = stocks[e["stock"]]
        d = e["day"]
        rows.append(dict(exec_id=e["exec_id"], order_id=by_pid[e["pid"]].order_id, isin=isin_of(s, d),
                         side=e["side"], last_qty=qty_view(s, d, e["qty"]), last_px="%.4f" % px_view(s, d, e["price"]),
                         last_mkt=venue_code(e["venue"], d), transact_time=K.iso_ms(e["utc"]), exec_type="FILL",
                         ref_exec_id="", trading_session="CONT" if e["session"] == "CONT" else "AUCT"))
        if e["busted"]:
            bt = e["utc"] + dt.timedelta(minutes=int(K.rng("bust", e["exec_id"]).integers(12, 95)))
            busts.append(dict(exec_id=e["exec_id"] + "X", order_id=by_pid[e["pid"]].order_id, isin=isin_of(s, d),
                              side=e["side"], last_qty=qty_view(s, d, e["qty"]),
                              last_px="%.4f" % px_view(s, d, e["price"]), last_mkt=venue_code(e["venue"], d),
                              transact_time=K.iso_ms(bt), exec_type="BUST", ref_exec_id=e["exec_id"],
                              trading_session="CONT", _t=bt))
    for b in busts:
        rows.append(b)
    rows.sort(key=lambda r: (r["transact_time"], r["exec_id"]))
    for r in rows:
        r.pop("_t", None)
    return rows


PEMB_HEAD = ["Trade Date", "Exec Time", "Client Order", "Exec Ref", "SEDOL", "B/S", "Shares", "Price (GBP)", "Venue",
             "Status", "Orig Exec Ref"]


def pemb_sheets(X):
    stocks = X["stocks"]
    by_pid = {p.pid: p for p in X["parents"]}
    sheets = {}
    for e in X["execs"]:
        if e["broker"] != "PEMB":
            continue
        s = stocks[e["stock"]]
        d = e["day"]
        off = K.london_offset(e["utc"])
        loc = e["utc"] + dt.timedelta(hours=off)
        key = (d.year, d.month)
        px = px_view(s, d, e["price"]) / 100.0
        row = [d.strftime("%d/%m/%Y"), loc.strftime("%H:%M:%S.") + "%03d" % (loc.microsecond // 1000),
               by_pid[e["pid"]].order_id, e["exec_id"], sedol_of(s, d), "B" if e["side"] == "B" else "S",
               qty_view(s, d, e["qty"]), round(px, 6), K.PEMBERWICK_VENUE[e["venue"]] if e["session"] == "CONT"
               else "WEXCOMBE CLOSE", "Executed", ""]
        sheets.setdefault(key, []).append((loc, row))
        if e["busted"]:
            bt = loc + dt.timedelta(minutes=int(K.rng("bust", e["exec_id"]).integers(12, 95)))
            row2 = [d.strftime("%d/%m/%Y"), bt.strftime("%H:%M:%S.") + "%03d" % (bt.microsecond // 1000),
                    by_pid[e["pid"]].order_id, e["exec_id"] + "-B", sedol_of(s, d), row[5], row[6], row[7], row[8],
                    "Busted", e["exec_id"]]
            sheets[key].append((bt, row2))
    return sheets


def write_pemb(path, X):
    wb = Workbook(write_only=True)
    for (y, m), items in sorted(pemb_sheets(X).items()):
        ws = wb.create_sheet("%s %d" % (K.MON3[m], y))
        ws.append(PEMB_HEAD)
        for _, r in sorted(items, key=lambda x: (x[0], x[1][3])):
            ws.append(r)
    wb.save(path)


OMS_COLS = ["order_id", "fund", "isin", "side", "quantity", "released_utc", "broker", "strategy", "arrival_mid_gbx",
            "status", "completed_utc"]


def oms_rows(X):
    stocks = X["stocks"]
    rows = []
    for p in X["parents"]:
        s = p.stock
        d = p.day
        rel = K.utc_of(d, p.t0)
        last = max(c[0] for c in p.children)
        done_t = K.UNCROSS_S if p.auction else last
        rows.append(dict(order_id=p.order_id, fund=p.fund, isin=isin_of(s, d), side="BUY" if p.side == "B" else "SELL",
                         quantity=qty_view(s, d, p.qty), released_utc=K.iso_ms(rel),
                         broker="ASHL" if (p.algo == "DMA" or K.ALGO[p.algo]["broker"] == "ASHL") else "PEMB",
                         strategy="DMA" if p.algo == "DMA" else K.ALGO[p.algo]["name"],
                         arrival_mid_gbx="%.4f" % px_view(s, d, p.arrival_mid), status="FILLED",
                         completed_utc=K.iso_ms(K.utc_of(d, done_t))))
    rows.sort(key=lambda r: (r["released_utc"], r["order_id"]))
    return rows


def write_venues(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Venues"
    head = ["MIC", "Venue", "Operator", "Type", "Minimum order", "Pemberwick name", "Valid from", "Valid to", "Notes"]
    ws.append(head)
    for v in K.VENUES:
        info = K.VENUE[v]
        if v == "THGD":
            ws.append([v, info["name"], info["operator"], info["kind"], "none", K.PEMBERWICK_VENUE[v], "2023-03-01",
                       "2026-05-29", "Segment code withdrawn when the operator moved the book to its new matching "
                       "platform"])
            ws.append([K.THURSGATE_NEW, info["name"], info["operator"], info["kind"], "none", K.PEMBERWICK_VENUE[v],
                       "2026-06-01", "", "Replaces THGD"])
        else:
            ws.append([v, info["name"], info["operator"], info["kind"], "none", K.PEMBERWICK_VENUE[v],
                       "2021-01-04" if v in ("XWEX", "AWKL", "DLVL") else "2022-06-06", "",
                       "Closing auction on this book" if v == "XWEX" else ""])
    for c, w in zip("ABCDEFGHI", (8, 30, 30, 24, 14, 18, 12, 12, 70)):
        ws.column_dimensions[c].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    wb.save(path)


MASTER_COLS = ["isin", "sedol", "ticker", "name", "sector", "index_band", "currency", "adv_shares_2026", "tick_gbx"]


def master_rows(X):
    rows = []
    for s in X["stocks"]:
        rows.append(dict(isin=s["isin"], sedol=s["sedol"], ticker=s["ticker"], name=s["name"], sector=s["sector"],
                         index_band=s["band"], currency="GBX", adv_shares_2026=int(round(s["adv"])),
                         tick_gbx="%g" % W.tick(s["price"])))
    rows.sort(key=lambda r: r["ticker"])
    return rows


def write_ca(path, X):
    c = X["stocks"][W.CONSOLIDATED_INDEX]
    other = X["stocks"][12]
    other2 = X["stocks"][88]
    events = [
        dict(event_id="CA-2026-0141", event_type="SPECIAL_DIVIDEND", isin=other["isin"], name=other["name"],
             ex_date="2026-03-12", pay_date="2026-04-17", amount_gbx="21.5000", ratio=None, new_isin=None,
             new_sedol=None, new_name=None),
        dict(event_id="CA-2026-0217", event_type="CONSOLIDATION", isin=c["old_isin"], name=c["name"],
             ex_date=W.CONSOLIDATION_DATE.isoformat(), pay_date=None, amount_gbx=None, ratio="1:10",
             new_isin=c["isin"], new_sedol=c["sedol"], new_name=None),
        dict(event_id="CA-2026-0388", event_type="NAME_CHANGE", isin=other2["isin"], name=W.FORMER_NAME[88],
             ex_date="2026-08-03", pay_date=None, amount_gbx=None, ratio=None, new_isin=None, new_sedol=None,
             new_name=other2["name"]),
    ]
    doc = dict(source="Fenmarsh Investment Management, securities reference team", extracted="2026-10-30",
               market="UK", events=events)
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        json.dump(doc, fh, indent=2)
        fh.write("\n")


CAL_COLS = ["date", "weekday", "uk_market", "note"]
HOLIDAY_NAMES = {dt.date(2025, 12, 25): "Christmas Day", dt.date(2025, 12, 26): "Boxing Day",
                 dt.date(2026, 1, 1): "New Year's Day", dt.date(2026, 4, 3): "Good Friday",
                 dt.date(2026, 4, 6): "Easter Monday", dt.date(2026, 5, 4): "Early May bank holiday",
                 dt.date(2026, 5, 25): "Spring bank holiday", dt.date(2026, 8, 31): "Summer bank holiday",
                 dt.date(2026, 12, 25): "Christmas Day", dt.date(2026, 12, 28): "Boxing Day (substitute day)",
                 dt.date(2027, 1, 1): "New Year's Day"}


def cal_rows():
    rows = []
    d = dt.date(2025, 12, 1)
    while d <= dt.date(2027, 1, 31):
        closed = d.weekday() >= 5 or d in K.BANK_HOLIDAYS
        note = HOLIDAY_NAMES.get(d, "")
        if d in (dt.date(2025, 12, 24), dt.date(2025, 12, 31), dt.date(2026, 12, 24), dt.date(2026, 12, 31)):
            note = "Early close 12:30"
        rows.append(dict(date=d.isoformat(), weekday=d.strftime("%a"), uk_market="CLOSED" if closed else "OPEN",
                         note=note))
        d += dt.timedelta(days=1)
    return rows


ORDER_COLS = ["client_order_ref", "account_ref", "wrapper", "isin", "side", "basis", "shares", "value_gbp",
              "placed_at", "status"]


def order_rows(X):
    stocks = X["stocks"]
    rows = []
    for o in X["orders"]:
        s = stocks[o["stock"]]
        pday = o["point_day"]
        status = o["status"]
        if o["g_day"] > K.DATA_END and status == "DEALT":
            status = "PENDING"
        shares = o["shares_req"] if o["basis"] == "SHARES" else ""
        rows.append(dict(client_order_ref=o["ref"], account_ref=o["account"], wrapper=o["wrapper"],
                         isin=s["isin"], side=o["side"], basis=o["basis"], shares=shares,
                         value_gbp=("%.2f" % o["value"]) if o["basis"] == "VALUE" else "",
                         placed_at=o["t_utc"].strftime("%Y-%m-%d %H:%M:%S"), status=status))
    return rows


NOTE_COLS = ["note_ref", "client_order_ref", "trade_date", "trade_time", "isin", "side", "shares", "price_gbx",
             "consideration_gbp", "commission_gbp", "ptm_levy_gbp", "stamp_duty_gbp", "settlement_date",
             "note_status"]


def old_amount(side, gross, stamp):
    """The old contract note system's amount: the market settlement of the trade, stamp duty in on a purchase."""
    return round(gross + stamp, 2) if side == "BUY" else round(gross, 2)


def net_amount(side, gross, comm, ptm, stamp):
    """The new contract note system's amount: what the client paid on a buy, what the client received on a sell."""
    return round(gross + comm + ptm + stamp, 2) if side == "BUY" else round(gross - comm - ptm, 2)


def settle(day):
    d = day
    for _ in range(2):
        d = K.next_business_day(d)
    return d


def note_rows(X):
    """Gresmere's notes as its back office extracts them: by settlement date, the old system's references before
    the 1 July migration and the new system's after, with the end-of-June trades issued by both."""
    stocks = X["stocks"]
    rg = K.rng("notes")
    rows = []
    n_old = n_new = 0
    for nt in sorted(X["notes"], key=lambda x: (x["day"], x["tmin"], x["seq"])):
        s = stocks[nt["stock"]]
        sd = settle(nt["day"])
        if sd < dt.date(2026, 1, 2) or sd > dt.date(2026, 11, 3):
            continue
        base = dict(client_order_ref=nt["order"], trade_date=nt["day"].strftime("%d/%m/%Y"),
                    trade_time="%02d:%02d" % (nt["tmin"] // 60, nt["tmin"] % 60), isin=isin_of(s, nt["day"]),
                    side=nt["side"], shares=nt["shares"], price_gbx="%.4f" % nt["price"],
                    consideration_gbp="%.2f" % old_amount(nt["side"], nt["consideration"], nt["stamp"]),
                    commission_gbp="%.2f" % nt["commission"],
                    ptm_levy_gbp="%.2f" % nt["ptm"], stamp_duty_gbp="%.2f" % nt["stamp"],
                    settlement_date=sd.strftime("%d/%m/%Y"), note_status="ISSUED")
        if nt["day"] < K.GRESMERE_MIGRATION:
            n_old += 1
            ref = "GSL%07d" % (4100000 + n_old)
        else:
            n_new += 1
            ref = "CN26-%08d" % (10000000 + n_new * 3)
            base["consideration_gbp"] = "%.2f" % net_amount(nt["side"], nt["consideration"], nt["commission"],
                                                             nt["ptm"], nt["stamp"])
        r = dict(base, note_ref=ref)
        reissue = rg.random() < 0.005
        if reissue:
            r["note_status"] = "CANCELLED"
        rows.append(r)
        if reissue:
            px2 = round(nt["price"] * (1 + rg.normal(0, 0.0015)), 4)
            cons2 = round(nt["shares"] * px2 / 100.0, 2)
            if nt["day"] < K.GRESMERE_MIGRATION:
                n_old += 1
                ref2 = "GSL%07d" % (4100000 + n_old)
            else:
                n_new += 1
                ref2 = "CN26-%08d" % (10000000 + n_new * 3)
            sd2 = round(cons2 * 0.005, 2) if nt["side"] == "BUY" else 0.0
            shown = old_amount(nt["side"], cons2, sd2) if nt["day"] < K.GRESMERE_MIGRATION else \
                net_amount(nt["side"], cons2, nt["commission"], nt["ptm"], sd2)
            rows.append(dict(base, note_ref=ref2, price_gbx="%.4f" % px2, consideration_gbp="%.2f" % shown,
                             stamp_duty_gbp="%.2f" % sd2, note_status="REISSUED"))
            nt["reissued_cons"] = cons2
        # the new system issued its own note for every trade settling on or after the migration date that the old
        # system had already issued (29 and 30 June), at its two-decimal price
        if nt["day"] < K.GRESMERE_MIGRATION <= sd and not reissue:
            n_new += 1
            p2 = round(nt["price"], 2)
            c2 = round(nt["shares"] * p2 / 100.0, 2)
            sd3 = round(c2 * 0.005, 2) if nt["side"] == "BUY" else 0.0
            rows.append(dict(base, note_ref="CN26-%08d" % (10000000 + n_new * 3), price_gbx="%.4f" % p2,
                             consideration_gbp="%.2f" % net_amount(nt["side"], c2, nt["commission"], nt["ptm"], sd3),
                             stamp_duty_gbp="%.2f" % sd3))
    return rows


AUCT_COLS = ["date", "isin", "uncrossing_price_gbx", "auction_volume_shares", "auction_value_gbp"]


def auction_rows(X):
    rg = K.rng("auction")
    rows = []
    for d in X["days"]:
        for s in X["stocks"]:
            close = X["closes"][s["i"]][d]
            vol = int(s["adv"] * float(rg.uniform(0.06, 0.14)))
            px = px_view(s, d, close)
            q = qty_view(s, d, vol)
            rows.append(dict(date=d.isoformat(), isin=isin_of(s, d), uncrossing_price_gbx="%.4f" % px,
                             auction_volume_shares=q, auction_value_gbp="%.2f" % (q * px / 100.0)))
    return rows
