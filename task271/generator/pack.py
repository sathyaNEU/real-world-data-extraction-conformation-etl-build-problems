"""Read the pack as a solver would: the survey's sample, the two centres' call records and the interview file conformed
into one call history per sampled member; the roster, the change log and the household extract for the planning
figures. Everything here reads the shipped files only."""
import csv
import datetime as dt
import os

from openpyxl import load_workbook

import common as K

DAY_END = dt.time(16, 30)
EVE_START = dt.time(17, 30)
REACHED_CRM = {"Spoke to member - survey completed": 1, "Spoke to member - declined survey": 2}
REACHED_DIALER = {"1100": 1, "2110": 2}


def _csv(path, delimiter=","):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f, delimiter=delimiter))


def regions(target):
    wb = load_workbook(os.path.join(target, K.FILES["regions"]), read_only=True)
    county = {}
    for row in list(wb["Counties"].iter_rows(values_only=True))[1:]:
        county[row[0]] = row[1]
    return county


def window(when):
    """0 weekday day, 1 weekday evening, 2 Saturday, from the local (Pacific) time of the call."""
    if when.weekday() == 5:
        return 2
    return 1 if when.time() >= EVE_START else 0


def _local(stamp):
    """An ISO time with its offset from UTC, as Pacific local time (daylight time through the field period)."""
    t = dt.datetime.fromisoformat(stamp)
    return (t.astimezone(dt.timezone(dt.timedelta(hours=-7)))).replace(tzinfo=None)


def _ampm(d, s):
    t = dt.datetime.strptime(s, "%I:%M %p").time()
    return dt.datetime.combine(d, t)


def survey(target):
    """Per sampled case: region, age, sex, release, the calls in order (time, centre, window, outcome code 0, 1 or 2)
    and the class where the interview file gives it (1 worked 80 hours or more in March, 0 not)."""
    sample = {r["case_id"]: r for r in _csv(os.path.join(target, K.FILES["sample"]))}
    by_member = {r["member_id"]: cid for cid, r in sample.items()}
    calls = {cid: [] for cid in sample}
    for r in _csv(os.path.join(target, K.FILES["dialer"])):
        if r["case_id"] not in sample:
            continue                                   # the February pilot and the test calls of 10 April
        when = _local(r["call_start"])
        calls[r["case_id"]].append((when, 1, window(when), REACHED_DIALER.get(r["disposition"], 0)))
    wb = load_workbook(os.path.join(target, K.FILES["crm"]), read_only=True)
    rows = list(wb["Activity"].iter_rows(values_only=True))
    head = rows[0]
    for row in rows[1:]:
        rec = dict(zip(head, row))
        if rec["Activity Type"] != "MWIS outbound call":
            continue
        d = rec["Activity Date"]
        d = d.date() if isinstance(d, dt.datetime) else d
        when = _ampm(d, rec["Activity Time"])
        calls[by_member[rec["Member ID"]]].append((when, 0, 0, REACHED_CRM.get(rec["Result"], 0)))
    cls = {}
    for r in _csv(os.path.join(target, K.FILES["interviews"])):
        if r["case_id"] not in sample:
            continue                                   # the test interviews
        if r["status"] == "Complete":
            cls[r["case_id"]] = int(int(r["total_hours_march"]) >= 80)
        else:
            cls[r["case_id"]] = int(r["key_80_hours_march"] == "Yes")
    out = {}
    for cid, s in sample.items():
        cl = sorted(calls[cid], key=lambda c: (c[0], c[1]))
        out[cid] = dict(region=s["region"], age=int(s["age"]), sex=s["sex"], release=s["release"],
                        first=s["first_contact_by"], calls=cl, cls=cls.get(cid, -1))
    return out


def band(age):
    return 0 if age <= 29 else (1 if age <= 44 else 2)


def histories(S):
    """The model's input: per case (region index, age band, sex, class or -1, outcome, [(combo, attempt), ...]) with
    combo 0 a member services call, 1 to 3 a vendor call by day, evening or Saturday, up to the reaching call."""
    hist = []
    for cid in sorted(S):
        s = S[cid]
        seq, out = [], 0
        for k, (when, ctr, win, res) in enumerate(s["calls"], 1):
            seq.append((0 if ctr == 0 else 1 + win, k))
            if res:
                out = res
                break
        hist.append((K.CODES.index(s["region"]), band(s["age"]), 0 if s["sex"] == "F" else 1, s["cls"], out, seq))
    return hist


def _date(s, fmt):
    return dt.datetime.strptime(s, fmt).date()


def enrollment(target):
    """The roster (19 October), the change log (1 to 19 October) and the household extract (15 October)."""
    roster = _csv(os.path.join(target, K.FILES["roster"]))
    log = _csv(os.path.join(target, K.FILES["changes"]), delimiter="|")
    cases = _csv(os.path.join(target, K.FILES["cases"]))
    for r in roster:
        r["born"] = _date(r["birth_date"], "%m/%d/%Y")
        r["begin"] = _date(r["enrollment_begin"], "%Y-%m-%d")
        r["end"] = _date(r["scheduled_end"], "%Y-%m-%d") if r["scheduled_end"] else None
    for r in log:
        r["born"] = _date(r["birth_date"], "%Y%m%d")
        r["eff"] = _date(r["effective_date"], "%Y%m%d")
        r["at"] = dt.datetime.strptime(r["logged_at"], "%Y-%m-%d %H:%M:%S")
    for r in cases:
        r["born"] = _date(r["birth_date"], "%m/%d/%Y")
        r["begin"] = _date(r["member_begin"], "%Y-%m-%d")
        r["end"] = _date(r["member_end"], "%Y-%m-%d") if r["member_end"] else None
    return roster, log, cases
