"""The golden deliverable from the shipped pack alone, ellsmere_aggravated_assault.xlsx, and generator/figures.json,
which carries every figure the write-up states."""
import json
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

import methods as M
import pack as P

XLSX = "ellsmere_aggravated_assault.xlsx"
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def load(target):
    f = lambda k: os.path.join(target, P.FILES[k])
    rows = M.read_csv(f("extract"))
    ra = {}
    for y in (2021, 2022, 2023, 2024):
        ra.update(M.read_return_a(f("ra_%d" % y)))
    return {"rows": rows, "ra": ra, "dvu": M.read_csv(f("dvu")), "gvr": M.read_csv(f("gvr")), "me": M.read_csv(f("me"))}


def compute(d):
    rows, ra, dvu = d["rows"], d["ra"], d["dvu"]
    lad = M.ladder(rows, ra, dvu)
    by_month, rate, bridge = M.restate(rows, ra, dvu)
    n13 = M.nibrs_13a(rows)
    tot, to13a = M.dual_tables(rows)
    dv_js = M.dvu_4e(dvu)
    e4_js = sum(ra[ym]["4E"][2] for ym in M.JAN_SEP)
    dv_q4 = tot["4E_dv"] / (tot["4E_dv"] + tot["4E_non"])
    published = {ym: sum(ra[ym][l][2] for l in M.AGG) for ym in M.MONTHS_2024}
    pub24 = sum(published[ym] for ym in M.MONTHS_2024)
    rest24 = sum(by_month.values())
    y25 = sum(n13[ym] for ym in M.MONTHS_2025)
    vic, sho = M.shootings(d["gvr"], d["me"])
    rounded = {"%d-%02d" % ym: round(v) for ym, v in by_month.items()}
    fig = {
        "call": {"change_pct": int(round(100 * (y25 / rest24 - 1))), "change_exact": 100 * (y25 / rest24 - 1)},
        "published_2024": pub24, "restated_2024": int(round(rest24, -2)), "restated_2024_units": round(rest24),
        "restated_2024_exact": rest24, "nibrs_2025": y25,
        "published_change_pct": round(100 * (y25 / pub24 - 1), 1),
        "bridge": {"hidden": round(bridge["hidden"]), "gained": round(bridge["gained"]), "lost": round(bridge["lost"]),
                   "hidden_exact": bridge["hidden"], "gained_exact": bridge["gained"], "lost_exact": bridge["lost"]},
        "restated_by_month": rounded,
        "restated_by_month_exact": {"%d-%02d" % ym: v for ym, v in by_month.items()},
        "published_by_month": {"%d-%02d" % ym: published[ym] for ym in M.MONTHS_2024},
        "nibrs_2025_by_month": {"%d-%02d" % ym: n13[ym] for ym in M.MONTHS_2025},
        "rates": {k: rate[k] for k in sorted(rate)},
        "dual": {"tot": dict(tot), "to13a": dict(to13a), "dv_share_q4": dv_q4,
                 "dv_share_jan_sep": sum(dv_js.values()) / e4_js, "dv_4e_jan_sep": sum(dv_js.values()),
                 "e4_jan_sep": e4_js},
        "routes": {k: (100 * v if isinstance(v, float) and abs(v) < 1 else v) for k, v in lad.items()},
        "shootings": {"victims": vic, "shootings": sho},
    }
    # every reasonable grain of the conversion, at the precision the figures are given
    fig["grains"] = {}
    for name, kw in M.GRAINS.items():
        t = M.restate_grain(rows, ra, dvu, **kw)
        fig["grains"][name] = {"restated_2024": t, "change_pct": 100 * (y25 / t - 1)}
    fig["routes"]["per_line_2024_hundred"] = int(round(lad["per_line_2024"], -2))
    fig["routes"]["per_line_pct"] = int(round(100 * lad["per_line"]))
    fig["routes"]["per_line_fall_pct"] = abs(fig["routes"]["per_line_pct"])
    # the shares and rates as the write-up prints them
    fig["printed"] = {"dv_share_q4_pct": round(100 * fig["dual"]["dv_share_q4"], 1),
                      "dv_share_jan_sep_pct": round(100 * fig["dual"]["dv_share_jan_sep"], 1),
                      "rate_4E_dv_pct": round(100 * rate["4E_dv"], 1), "rate_4E_non_pct": round(100 * rate["4E_non"], 2),
                      "rate_4D_pct": round(100 * rate["4D"], 1), "rate_4A_pct": round(100 * rate["4A"], 1)}
    return fig


def write_xlsx(path, fig):
    wb = openpyxl.Workbook()
    head = PatternFill("solid", fgColor="D9E1EA")
    bold = Font(bold=True)
    small = Font(italic=True, size=9, color="595959")
    count, signed, pct = "#,##0", "#,##0;-#,##0", "+0%;-0%"
    ws = wb.active
    ws.title = "Aggravated assault"
    ch = fig["call"]["change_pct"]
    ws["A1"] = "Aggravated assault %s %d%% in 2025 once 2024 is counted the way NIBRS counts it" % (
        "rose" if ch > 0 else "fell", abs(ch))
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Ellsmere Police Department. Victims of NIBRS offense 13A, one offense per victim, by the month the "
                "offense was reported (Ellsmere local time). 2024 restated as NIBRS would have recorded it had the "
                "department reported to NIBRS all year; October to December 2024 counted from the extract.")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:D2")
    ws.row_dimensions[2].height = 44
    for j, h in enumerate(("", "2024 restated", "2025", "Change")):
        c = ws.cell(row=4, column=j + 1, value=h or None)
        c.font, c.fill = bold, head
        c.alignment = Alignment(horizontal="left" if j == 0 else "right")
    ws.cell(row=5, column=1, value="Aggravated assault victims (NIBRS 13A)").font = bold
    for j, (v, fmt) in enumerate(((fig["restated_2024"], count), (fig["nibrs_2025"], count),
                                  (round(ch / 100, 2), pct))):
        c = ws.cell(row=5, column=2 + j, value=v)
        c.number_format, c.font = fmt, bold
    ws.cell(row=7, column=1, value="From the 2024 published figure to the restated one").font = bold
    for j, h in enumerate(("Line", "Victims")):
        c = ws.cell(row=8, column=j + 1, value=h)
        c.font, c.fill = bold, head
        c.alignment = Alignment(horizontal="left" if j == 0 else "right")
    walk = [("2024 as published (Return A, lines 4a to 4d, actual offenses)", fig["published_2024"]),
            ("Assault victims Return A did not score (incident scored as robbery or homicide under the hierarchy "
             "rule)", fig["bridge"]["hidden"]),
            ("Victims Return A scored as other assaults (4e) that NIBRS counts as aggravated", fig["bridge"]["gained"]),
            ("Victims Return A scored as aggravated (4a to 4d) that NIBRS counts as simple assault or intimidation",
             -fig["bridge"]["lost"]),
            ("2024 restated on the NIBRS basis", fig["restated_2024_units"])]
    for i, (k, v) in enumerate(walk):
        ws.cell(row=9 + i, column=1, value=k)
        ws.cell(row=9 + i, column=2, value=v).number_format = signed
    ws.cell(row=13, column=1).font = bold
    ws.cell(row=13, column=2).font = bold
    ws.cell(row=14, column=1, value="The restated total is an estimate, given to the nearest hundred above and to the "
                                    "unit here; the three middle lines are rounded from unrounded conversions. "
                                    "Sources: Return A 2024, column 4; NIBRS offense extract run 5 October 2026; "
                                    "Domestic Violence Unit case log exported 2 October 2026.").font = small
    ws.cell(row=14, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A14:D14")
    ws.row_dimensions[14].height = 30
    ws.column_dimensions["A"].width = 84
    for col in "BCD":
        ws.column_dimensions[col].width = 14
    ws2 = wb.create_sheet("2024 by month")
    for j, h in enumerate(("Month", "Published (Return A)", "Restated (NIBRS basis)", "Basis")):
        c = ws2.cell(row=1, column=j + 1, value=h)
        c.font, c.fill = bold, head
    for i, m in enumerate(range(1, 13)):
        key = "2024-%02d" % m
        ws2.cell(row=2 + i, column=1, value="%s 2024" % MONTHS[m - 1])
        ws2.cell(row=2 + i, column=2, value=fig["published_by_month"][key]).number_format = count
        ws2.cell(row=2 + i, column=3, value=fig["restated_by_month"][key]).number_format = count
        ws2.cell(row=2 + i, column=4, value="counted from the extract" if m >= 10 else "converted from Return A")
    ws2.cell(row=14, column=1, value="Total").font = bold
    for j, v in ((2, fig["published_2024"]), (3, fig["restated_2024_units"])):
        c = ws2.cell(row=14, column=j, value=v)
        c.font, c.number_format = bold, count
    if sum(fig["restated_by_month"].values()) != fig["restated_2024_units"]:
        ws2.cell(row=15, column=1, value="Converted months are rounded one by one, so they do not add exactly to the "
                                         "total.").font = small
    for col, w in zip("ABCD", (12, 21, 23, 26)):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = "A2"
    ws3 = wb.create_sheet("Shootings")
    for j, h in enumerate(("Quarter", "Non-fatal shooting victims", "Shootings they were hurt in")):
        c = ws3.cell(row=1, column=j + 1, value=h)
        c.font, c.fill = bold, head
    for i, q in enumerate(M.QUARTERS):
        ws3.cell(row=2 + i, column=1, value=q.replace("Q", " Q"))
        ws3.cell(row=2 + i, column=2, value=fig["shootings"]["victims"][q]).number_format = count
        ws3.cell(row=2 + i, column=3, value=fig["shootings"]["shootings"][q]).number_format = count
    ws3.cell(row=15, column=1, value="Sources: Gun Violence Reduction Unit shooting log, 1 October 2026; Medical "
                                     "Examiner homicide list, 6 October 2026.").font = small
    for col, w in zip("ABC", (10, 26, 28)):
        ws3.column_dimensions[col].width = w
    ws3.freeze_panes = "A2"
    ws4 = wb.create_sheet("Notes")
    notes = [
        "Prepared for the crime chapter of the 2026 Public Safety Scorecard. Counts are as the department's records "
        "stood when the extract was run on 5 October 2026.",
        "October to December 2024: every victim carries both its Return A line and its NIBRS offense. Rates from those "
        "months: aggravated lines 4a to 4d each at its own share counted 13A; robbery and homicide incidents at the "
        "13A victims per incident the hierarchy left unscored; other assaults (4e) separately for domestic victims "
        "(%.1f%% counted 13A) and the rest (%.2f%%)." % (fig["printed"]["rate_4E_dv_pct"],
                                                         fig["printed"]["rate_4E_non_pct"]),
        "January to September 2024: each month's Return A actual offenses at those rates, the domestic victims of line "
        "4e taken from the Domestic Violence Unit's log (cases not unfounded). Domestic victims were %.1f%% of 4e in "
        "October to December and %.1f%% in January to September." % (fig["printed"]["dv_share_q4_pct"],
                                                                      fig["printed"]["dv_share_jan_sep_pct"]),
        "Shootings: people struck in an assault within the city's limits who survived, each once, by the quarter shot "
        "in local time; a shooting is the people struck at one place and time. The shooting log's times are UTC on "
        "the rows the records system entered from October 2024 and are put in local time first. Deaths after the "
        "log's outcome was entered are taken from the Medical Examiner's list. The North Ellsmere Township blocks, "
        "which the log marked inside the city until September 2024, are left out.",
    ]
    for i, t in enumerate(notes):
        c = ws4.cell(row=1 + i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws4.row_dimensions[1 + i].height = 46
    ws4.column_dimensions["A"].width = 110
    wb.save(path)


def main(target, gdir, fig_path):
    os.makedirs(gdir, exist_ok=True)
    fig = compute(load(target))
    write_xlsx(os.path.join(gdir, XLSX), fig)
    import postprocess as PP
    PP.normalize_ooxml(os.path.join(gdir, XLSX), "Darlene Ford", "Darlene Ford", "2026-10-07T14:05:00Z",
                       "2026-10-09T20:41:17Z", excel=True, company="University of Ellsmere")
    with open(fig_path, "w", encoding="utf-8") as fh:
        json.dump(fig, fh, indent=1, sort_keys=True)
    return fig
