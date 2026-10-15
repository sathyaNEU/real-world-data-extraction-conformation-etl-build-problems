"""The pack's documents and workbooks: the Capital Committee's siting charter, the supplier's power content label,
the agency's published rates for 2025 and 2024, the data dictionary and the extract manifest."""
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import LETTER                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import inch                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import pack as P    # noqa: E402
import world as W   # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("LiberationSerif", "Regular"), ("LiberationSerif-Bold", "Bold"), ("LiberationSerif-Italic", "Italic"),
              ("LiberationSans", "Regular"), ("LiberationSans-Bold", "Bold")):
    try:
        pdfmetrics.getFont(nm)
    except KeyError:
        pdfmetrics.registerFont(TTFont(nm, FONT_DIR + nm.split("-")[0] + "-" + f + ".ttf"))
INK = colors.HexColor("#1d1f21")
GREY = colors.HexColor("#5a5f66")
GREEN = colors.HexColor("#2f5d3a")
BODY = ParagraphStyle("body", fontName="LiberationSerif", fontSize=10.5, leading=14.5, textColor=INK, spaceAfter=6)
NUM = ParagraphStyle("num", parent=BODY, leftIndent=30, firstLineIndent=-30)
H1 = ParagraphStyle("h1", fontName="LiberationSans-Bold", fontSize=14, leading=18, textColor=GREEN, spaceAfter=2)
H1S = ParagraphStyle("h1s", fontName="LiberationSans", fontSize=10, leading=13, textColor=GREY, spaceAfter=10)
H2 = ParagraphStyle("h2", fontName="LiberationSans-Bold", fontSize=10.5, leading=14, textColor=INK, spaceBefore=7, spaceAfter=3)
CELL = ParagraphStyle("cell", fontName="LiberationSans", fontSize=8.8, leading=11, textColor=INK)
CELLB = ParagraphStyle("cellb", parent=CELL, fontName="LiberationSans-Bold")
SMALL = ParagraphStyle("small", fontName="LiberationSans", fontSize=8.5, leading=11, textColor=GREY)


def _doc(path, title, author, footer):
    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFont("LiberationSans", 7.5)
        canvas.setFillColor(GREY)
        canvas.drawString(1.0 * inch, 0.6 * inch, footer)
        canvas.drawRightString(LETTER[0] - 1.0 * inch, 0.6 * inch, "Page %d" % doc.page)
        canvas.restoreState()
    d = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=1.0 * inch, rightMargin=1.0 * inch, topMargin=0.9 * inch,
                          bottomMargin=0.95 * inch, title=title, author=author, subject="", creator="", producer="")
    return d, frame


def _table(rows, widths):
    data = [[Paragraph(str(c), CELLB if i == 0 else CELL) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa0a6")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eaed")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 4),
                           ("RIGHTPADDING", (0, 0), (-1, -1), 4)]))
    return t


# ---------------------------------------------------------------------------------------------- the charter
def write_charter(path):
    doc, frame = _doc(path, "2027 On-Site Solar Pilot: Siting Charter", W.ORG,
                      "%s  |  Capital Committee  |  2027 On-Site Solar Pilot: Siting Charter  |  Approved 15 September 2026"
                      % W.ORG)
    s = [Paragraph("2027 On-Site Solar Pilot: Siting Charter", H1),
         Paragraph("%s, Capital Committee. Approved 15 September 2026. Charter owner: %s, %s." %
                   (W.ORG, W.PEOPLE["director"], W.TITLES["director"]), H1S)]
    P_ = lambda n, t: s.append(Paragraph("<b>%s</b>&nbsp;&nbsp;%s" % (n, t), NUM))
    s.append(Paragraph("1. Purpose", H2))
    P_("1.1", "Corvane will install one rooftop solar array at one of its stores in 2027 as the pilot for a wider "
              "on-site generation program. The Committee names the pilot store in the 2027 capital plan, which it "
              "approves on 4 December 2026.")
    P_("1.2", "The pilot is sited to cut the carbon dioxide that the electric grid emits to serve Corvane.")
    s.append(Paragraph("2. Candidates", H2))
    P_("2.1", "The candidates are the eight stores in the pilot store sites file, each with the array that the 2026 "
              "feasibility study designed for its roof.")
    P_("2.2", "%s has modelled each array hour by hour on 2025 weather at its store. For this charter an array's "
              "2025 output is that model's output." % W.VENDOR_MODEL)
    s.append(Paragraph("3. The siting measure", H2))
    P_("3.1", "The pilot goes to the candidate whose array would have avoided the most grid carbon dioxide per "
              "megawatt-hour it generated in 2025, the latest calendar year of complete grid records.")
    P_("3.2", "The grid carbon dioxide an array avoids is the carbon dioxide that the generators of its store's balancing "
              "area would not have emitted in 2025 had the array been supplying the store.")
    P_("3.3", "Each megawatt-hour the array generates is one megawatt-hour less that the area's generators produce. The "
              "area's interchange with its neighbours is scheduled a day ahead and does not change with the array's "
              "output, and the store uses all of the array's output: each array is sized below its store's daytime load.")
    P_("3.4", "The measure is stated in metric tonnes of carbon dioxide per megawatt-hour of array output.")
    s.append(Paragraph("4. The recommendation", H2))
    P_("4.1", "The Director of Energy and Sustainability recommends the pilot store to the Committee in a memo that "
              "states the measure for every candidate and names the first candidate left out.")
    P_("4.2", "For the finance team's valuation of the array's output, the memo also states each candidate's 2025 metered "
              "electricity use, the share of it that its array would have supplied, and what Corvane paid for the "
              "store's electricity per megawatt-hour in 2025.")
    s.append(Paragraph("5. Records", H2))
    P_("5.1", "The %s's hourly operating data for each balancing area, its plant register and its published 2025 "
              "generation and emission rates; the array model; and Corvane's accounts payable records of its "
              "electricity invoices." % W.AGENCY)
    s.append(Spacer(1, 10))
    s.append(Paragraph("Approved by the Capital Committee at its meeting of 15 September 2026 (minute CC-26-41).", SMALL))
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------------------------------------- the supplier's label
def write_label(path):
    doc, frame = _doc(path, "Power Content Label 2025", W.SUPPLIER,
                      "%s  |  Power Content Label, calendar year 2025  |  Issued 30 June 2026" % W.SUPPLIER)
    s = [Paragraph("Power Content Label", H1),
         Paragraph("%s. Calendar year 2025. Issued 30 June 2026 under the retail electricity disclosure rules." % W.SUPPLIER,
                   H1S),
         Paragraph("This label discloses the sources of the electricity %s purchased in 2025 to serve customers on the "
                   "product below, and the carbon dioxide emitted generating it. It covers the product as sold in every "
                   "service territory where Ardley serves commercial customers." % W.SUPPLIER.replace(" LLC", ""), BODY),
         Paragraph("Product: Ardley Business Fixed (commercial fixed-price service)", H2)]
    rows = [["Energy source", "Share of product (%)", "Share of national grid mix (%)"],
            ["Coal", "19", "17"], ["Natural gas", "44", "41"], ["Nuclear", "18", "19"], ["Large hydroelectric", "6", "7"],
            ["Wind", "8", "10"], ["Solar", "3", "5"], ["Petroleum", "under 1", "under 1"], ["Unspecified sources", "2", "n/a"],
            ["Total", "100", "100"]]
    s.append(_table(rows, [2.6 * inch, 1.6 * inch, 1.9 * inch]))
    s.append(Spacer(1, 8))
    s.append(Paragraph("Carbon dioxide intensity of the product: 826 pounds per megawatt-hour delivered. Renewable energy "
                       "certificates retired for the product: none beyond the state renewable portfolio standard.", BODY))
    s.append(Paragraph("Sources are reported from supply contracts and wholesale purchases; unspecified sources are market "
                       "purchases not traceable to a generator. Figures are rounded and may not sum to 100.", SMALL))
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------------------------------------- workbooks
def write_rates(path, table, year, release):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BA%02d" % (year % 100)
    ws["A1"] = "%s, %s: data year %d. Release 1.0, %s." % (W.AGENCY, W.RATES_TITLE, year,
                                                             release.strftime("%-d %B %Y"))
    ws["A1"].font = Font(bold=True)
    hdr = ["BA code", "Balancing area name", "Annual net generation (MWh)", "Annual CO2 emissions (short tons)",
           "CO2 output emission rate (lb/MWh)", "CO2 fossil output emission rate (lb/MWh)",
           "CO2 coal output emission rate (lb/MWh)", "CO2 gas output emission rate (lb/MWh)",
           "CO2 oil output emission rate (lb/MWh)", "CO2 non-baseload output emission rate (lb/MWh)",
           "Coal generation (%)", "Gas generation (%)", "Oil generation (%)", "Nuclear generation (%)",
           "Hydro generation (%)", "Wind generation (%)", "Solar generation (%)"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=3, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCE6F1")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    r = 4
    for code in W.ORDER:
        t = table[code]
        vals = [code, W.AREAS[code]["ba"], int(t["gen"]), int(t["co2_short"]), t["avg"], t["fossil"], t["coal"], t["gas"],
                t["oil"], t["nb"]] + [round(100 * t["mix"][f], 1) for f in
                                      ("coal", "natural gas", "petroleum", "nuclear", "hydro", "wind", "solar")]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            if j in (3, 4):
                c.number_format = "#,##0"
            elif 5 <= j <= 10:
                c.number_format = "#,##0.0"
            elif j >= 11:
                c.number_format = "0.0"
        r += 1
    widths = [9, 38, 16, 16, 14, 14, 14, 14, 14, 16, 10, 10, 10, 10, 10, 10, 10]
    for j, wdt in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = wdt
    ws.row_dimensions[3].height = 60
    ws.freeze_panes = "C4"
    n = wb.create_sheet("Notes")
    notes = [
        "%s: %s" % (W.AGENCY, W.RATES_TITLE),
        "",
        "Coverage: every generating plant of 1 MW or more registered in the %s balancing areas, from the plant "
        "register for data year %d." % ("eight" if year else "", year),
        "Net generation: a plant's gross generation less its station service, as reported for the data year.",
        "CO2 emissions: from fuel consumed and the fuel's emission factor, in short tons (2,000 lb).",
        "Output emission rate: a balancing area's annual CO2 emissions divided by its annual net generation from all "
        "sources, in pounds per MWh.",
        "Fossil, coal, gas and oil output emission rates: the same ratio for the area's plants burning that fuel group.",
        "Non-baseload output emission rate: the output emission rate of the area's combustion plants, each plant weighted "
        "by its non-baseload share: wholly non-baseload with an annual capacity factor below 0.2, wholly baseload above "
        "0.8, and linear between.",
        "Generation shares: the area's annual net generation by primary fuel.",
        "Balancing areas are listed under the codes in use at the end of the data year.",
    ]
    for i, line in enumerate(notes, 1):
        if line:
            n.cell(row=i, column=1, value=line)
    n.column_dimensions["A"].width = 120
    wb.save(path)


def write_dictionary(path, counts):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    F = W.FILES
    rows = [
        (F["hourly"], "", "One row per balancing area and clock hour of 2025, the hour named by its end; MW are the hour's "
                          "average (so equal MWh in the hour). %s rows." % format(counts["hourly"], ",")),
        (F["hourly"], "balancing_area", "The area's code at the hour; an area that changed its code reports under its "
                                        "former code before the change (see %s)." % F["codes"]),
        (F["hourly"], "utc_time_at_end_of_hour", "End of the hour, UTC."),
        (F["hourly"], "local_time_at_end_of_hour", "End of the hour on the area's local clock, with its UTC offset."),
        (F["hourly"], "demand_forecast_mw", "The operator's day-ahead demand forecast for the hour."),
        (F["hourly"], "demand_mw", "Electricity consumed in the area plus its losses, as the operator reports it."),
        (F["hourly"], "demand_imputed", "Y where the hour's demand telemetry was lost and the operator filed demand "
                                         "from the hour's net generation and interchange, to the nearest 5 MW."),
        (F["hourly"], "net_generation_mw", "Net generation of the area's generators, all sources."),
        (F["hourly"], "total_interchange_mw", "Net flow to neighbouring areas; positive is export. Scheduled day-ahead."),
        (F["hourly"], "ng_coal_mw ... ng_other_mw", "Net generation by primary fuel; blank where the area has none. Solar "
                                                    "plants' station service shows as a small negative value at night."),
        (F["arrays"], "", "One row per store and hour of 2025 on the store's local standard time. %s rows." %
         format(counts["arrays"], ",")),
        (F["arrays"], "model_run", "%s model run for the store's designed array (2025 weather)." % W.VENDOR_MODEL),
        (F["arrays"], "hour_beginning_lst", "Start of the hour, local standard time at the store."),
        (F["arrays"], "ac_output_kw", "Modelled AC output at the inverter, average kW over the hour."),
        (F["plants"], "", "One row per plant registered in the eight areas, 2025."),
        (F["plants"], "net_generation_2025_mwh", "Annual net generation, MWh."),
        (F["plants"], "co2_2025_short_tons", "Annual CO2 emissions, short tons."),
        (F["plants"], "prime_mover", "ST steam turbine, CC combined cycle, GT combustion turbine, HY hydraulic turbine, "
                                     "WT wind turbine, PV photovoltaic."),
        (F["rates"], "", "%s, %s, data year 2025: one row per area (sheet BA25) and definitions (sheet Notes)." %
         (W.AGENCY, W.RATES_TITLE)),
        (F["rates24"], "", "The same publication for data year 2024."),
        (F["sites"], "", "One row per candidate store: its delivery utility and the array the feasibility study designed."),
        (F["codes"], "", "Balancing area codes, names and UTC offsets; a code that replaced another, from the date it took effect."),
        (F["utilmap"], "", "Electric utilities and the balancing area each serves."),
        (F["invoices"], "", "Corvane accounts payable: one row per line of each electricity invoice for the eight stores "
                            "with service in 2025. A vendor that bills a meter's service period again (on the actual "
                            "read after an estimate) issues a new invoice; the later invoice is the bill of record. %s "
                            "rows." % format(counts["invoices"], ",")),
        (F["invoices"], "service_period_start, service_period_end", "First and last day of service the invoice bills."),
        (F["invoices"], "quantity, unit_of_measure", "The quantity the line's charge is computed on, in its unit. "
                                                     "Delivery utilities bill the reading at the store's meter; Ardley "
                                                     "Retail Energy bills energy at the wholesale delivery point, the "
                                                     "meter's reading plus the delivery territory's losses."),
        (F["invoices"], "line_amount_usd", "The line's charge as billed, dollars; tax lines included."),
        (F["charter"], "", "The Capital Committee's charter for siting the 2027 pilot."),
        (F["label"], "", "Ardley Retail Energy's power content label for 2025."),
    ]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fields"
    hdr = ["File", "Field", "Definition"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E2EFDA")
    for i, r in enumerate(rows, 2):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v if v != "" else None)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 90
    ws.freeze_panes = "A2"
    wb.save(path)


def write_manifest(path, counts):
    F = W.FILES
    rows = [
        [F["hourly"], W.AGENCY + ", " + W.HOURLY_TITLE + ", bulk download", W.HOURLY_PULL.isoformat(), counts["hourly"],
         "csv", "public data"],
        [F["plants"], W.AGENCY + ", plant register, data year 2025", W.HOURLY_PULL.isoformat(), counts["plants"], "csv",
         "public data"],
        [F["rates"], W.AGENCY + ", " + W.RATES_TITLE + ", data year 2025, release 1.0", W.RATES_RELEASE.isoformat(), 8,
         "xlsx", "public data"],
        [F["rates24"], W.AGENCY + ", " + W.RATES_TITLE + ", data year 2024, release 1.0", W.RATES_2024_RELEASE.isoformat(),
         8, "xlsx", "public data"],
        [F["codes"], W.AGENCY + ", balancing area code list", W.HOURLY_PULL.isoformat(), counts["codes"], "csv",
         "public data"],
        [F["utilmap"], W.AGENCY + ", utility to balancing area table", W.HOURLY_PULL.isoformat(), counts["utilmap"], "csv",
         "public data"],
        [F["arrays"], W.VENDOR_MODEL + ", array output model runs for the eight candidate stores", W.MODEL_RUN.isoformat(),
         counts["arrays"], "csv", "Corvane contract deliverable"],
        [F["sites"], "Corvane real estate, pilot store sites and 2026 feasibility designs", W.PACK_DATE.isoformat(),
         counts["sites"], "csv", "Corvane internal"],
        [F["invoices"], "Corvane accounts payable, electricity invoice lines with service in 2025 (vendor invoices; "
                        "credit memos and payments not included)", W.AP_PULL.isoformat(), counts["invoices"], "csv",
         "Corvane internal"],
        [F["charter"], "Corvane Capital Committee, siting charter", W.CHARTER_DATE.isoformat(), "", "pdf",
         "Corvane internal"],
        [F["label"], W.SUPPLIER + ", power content label 2025", W.LABEL_DATE.isoformat(), "", "pdf", "public disclosure"],
        [F["dict"], "Corvane energy team, field definitions for this folder", W.PACK_DATE.isoformat(), "", "xlsx",
         "Corvane internal"],
    ]
    return P.write_csv(path, ["file_name", "contents", "as_of", "rows", "format", "terms"], rows)
