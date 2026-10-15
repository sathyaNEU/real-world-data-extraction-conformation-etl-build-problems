"""task460's generator checks: the input gates; the call, every rung, reflex and grid cell with its margin and
Gorsefield's place; the estimator family's convergence on every graded figure; the three rules executed (R1 the field's
raw material ships, R2 each Layer 1 item mishandled leaves the call and the critical components, R3a the easy routes on
the clean truth file their wrong stores with every filed rule applied, R3b the decisive step alone files the golden,
R3c nothing planted, no document naming a basis or an estimator); the semantic forks' ruling sentences; the asks'
figures, devices and cracker paths; the asks' separation from the call; the realism sweep; the goldens."""
import copy
import csv
import datetime as dt
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import arrays as A      # noqa: E402
import grid as Gd       # noqa: E402
import methods as M     # noqa: E402
import world as W       # noqa: E402

TASK = os.path.dirname(HERE)
COUNT, FAILS = [0], []
PICK, SECOND = "GRSF", "HLVN"
RUNG = {"avg": "BRKW", "nb": "QLBY", "pooled": "TNFD"}           # R0, R1, R2
CRITICAL = ("GRSF", "HLVN", "QLBY")                              # the pick, the first left out, component 3
QUOTES = [
    ("SF1", "charter", "The grid carbon dioxide an array avoids is the carbon dioxide that the generators of its store's "
                       "balancing area would not have emitted in 2025 had the array been supplying the store."),
    ("SF2", "charter", "The area's interchange with its neighbours is scheduled a day ahead and does not change with the "
                       "array's output"),
    ("SF3", "charter", "For this charter an array's 2025 output is that model's output."),
    ("SF4", "charter", "Each megawatt-hour the array generates is one megawatt-hour less that the area's generators produce."),
    ("SF5", "charter", "the store uses all of the array's output: each array is sized below its store's daytime load."),
    ("SF6", "charter", "The measure is stated in metric tonnes of carbon dioxide per megawatt-hour of array output."),
    ("SF7", "dictionary", "One row per store and hour of 2025 on the store's local standard time."),
    ("SF8", "charter", "per megawatt-hour it generated in 2025, the latest calendar year of complete grid records."),
    ("SF9", "dictionary", "End of the hour, UTC."),
    ("SF9", "dictionary", "Start of the hour, local standard time at the store."),
    ("SF10", "dictionary", "an area that changed its code reports under its former code before the change"),
    ("SF11", "charter", "states the measure for every candidate and names the first candidate left out."),
]
# words that would name a basis, an estimator or the decisive quantity; none may appear in the documents that govern
# the call
METHOD_WORDS = re.compile(r"marginal|margin\b|non-baseload|average|regress|slope|season|intensity|displace|"
                          r"operating margin|merit order|load-following|follow the load", re.I)
AVOID_WORDS = re.compile(r"avoid|displace|reduction|offset|marginal|margin\b|counterfactual", re.I)
DICT_WORDS = re.compile(r"marginal|margin\b|non-baseload|regress|slope|season|avoid|displace|intensity|"
                        r"average (?:output )?(?:emission )?rate|counterfactual", re.I)


def ok(cond, msg):
    COUNT[0] += 1
    if cond:
        print("ok    " + msg)
    else:
        print("FAIL  " + msg)
        FAILS.append(msg)


def text_of(path):
    if path.endswith(".pdf"):
        from pypdf import PdfReader
        return re.sub(r"\s+", " ", " ".join(p.extract_text() for p in PdfReader(path).pages))
    if path.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True)
        return " ".join(" ".join(str(c) for c in r if c is not None) for ws in wb.worksheets
                        for r in ws.iter_rows(values_only=True))
    return open(path, encoding="utf-8", errors="replace").read()


def edge(x, step):
    """distance from x to the nearest rounding edge when x is rounded to multiples of step"""
    f = (x / step) % 1.0
    return abs(f - 0.5) * step


def r2(x):
    return float("%.2f" % x)


def by_area(D, fig, key=None):
    return {D["stores"][st]["area"]: (fig[st] if key is None else fig[st][key]) for st in fig}


def order(vals):
    return sorted(vals, key=lambda c: -vals[c])


# ---------------------------------------------------------------------------------------------- the clean truth
def truth_routes(x):
    """Every route on the generator's clean truth: the dispatch before any rounding, renaming, imputation or clock
    convention, the agency's rates as the plants' true annual figures, each array's modelled output on the UTC
    clock. R3a takes the easy routes with every filed rule applied (per MWh of the array's output, the area's own
    generators, metric tonnes); R3b changes only the decisive step."""
    out = {}
    for code in W.ORDER:
        g = x["grid"][code]
        pub = Gd.published(g)
        fr = pub["rate_fuel"]
        E = sum(g["fuel"][f] * fr.get(f, 0.0) for f in g["fuel"])
        NG = sum(g["fuel"].values()) + g["nuc"] + g["hydro"] + g["wind"] + g["usolar"]
        se = np.array([M.SEASON[m] for m in g["mon"]])
        dE = np.diff(E, prepend=E[0])
        dN = np.diff(NG, prepend=NG[0])
        dD = np.diff(g["load"], prepend=g["load"][0])
        mask = np.ones(len(E), bool)
        mask[0] = False
        m = M.binned_slope(dE, dN, se * 24 + g["lh"], mask)
        x0, y0 = dD[1:] - dD[1:].mean(), dE[1:] - dE[1:].mean()
        idx = A.utc_index(code)
        S = x["arrays"][code]
        w = np.zeros(len(E))
        okk = idx >= 0
        np.add.at(w, idx[okk], S[okk])
        out[code] = dict(avg=pub["avg"], nb=pub["nb"], pooled=float((x0 * y0).sum() / (x0 * x0).sum()),
                         hourly_average=float(((E / NG) * w).sum() / w.sum()), all_hours=float(m.mean()),
                         decisive=float((m * w).sum() / w.sum()), true=float((g["mer"] * w).sum() / w.sum()),
                         true_all=float(g["mer"].mean()))
    return out


def copy_without(target, drop):
    tmp = tempfile.mkdtemp(dir=TASK, suffix=".chk")
    for f in os.listdir(target):
        if f not in drop:
            shutil.copy2(os.path.join(target, f), os.path.join(tmp, f))
    return tmp


# ---------------------------------------------------------------------------------------------- the checks
def run(target, golden_dir, x=None, meta_path=None, fig_path=None):
    COUNT[0] = 0
    del FAILS[:]
    files = sorted(os.listdir(target))
    F = W.FILES
    meta = json.load(open(meta_path)) if meta_path else None

    # ------------------------------------------------------------------ input gates
    fmts = {f.rsplit(".", 1)[1] for f in files}
    ok(len(files) >= 10, "input gate: %d files (10 or more)" % len(files))
    ok(len(fmts) >= 3, "input gate: %d formats %s (3 or more)" % (len(fmts), sorted(fmts)))
    rows = {f: sum(1 for _ in open(os.path.join(target, f), encoding="utf-8")) - 1 for f in files if f.endswith(".csv")}
    big = sorted(f for f in rows if rows[f] >= 25000)
    ok(len(big) >= 1, "input gate: %s with %s rows (25,000 or more)" % (big, [rows[f] for f in big]))
    if meta:
        ok(sorted(meta["distractor_files"]) == sorted(F[k] for k in W.DISTRACTORS)
           and all(d in files for d in meta["distractor_files"]),
           "input gate: two distractors named in metadata.json and shipped: %s" % meta["distractor_files"])
    ok(not any("distractor" in f.lower() for f in files), "no file name says distractor")

    # ------------------------------------------------------------------ the call and the ladder
    D = M.load(target)
    g = M.golden(D)
    R = M.rungs(D)
    fig = by_area(D, R)
    gold = {c: fig[c]["golden"] for c in fig}
    go = order(gold)
    ok(go[0] == PICK and go[1] == SECOND, "golden: %s first (%.4f), %s first left out (%.4f)" % (
        go[0], gold[go[0]], go[1], gold[go[1]]))
    ok(gold[PICK] / gold[SECOND] >= 1.5, "golden margin %.3f (1.5 or more)" % (gold[PICK] / gold[SECOND]))
    pos = {}
    for key, want in (("avg", "BRKW"), ("nb", "QLBY"), ("pooled", "TNFD")):
        v = {c: fig[c][key] for c in fig}
        o = order(v)
        pos[key] = o.index(PICK) + 1
        ok(o[0] == want and v[o[0]] / v[o[1]] >= 1.15,
           "rung %s: %s first at %.3f, %.3f over %s %.3f; Gorsefield %d" % (key, o[0], v[o[0]], v[o[0]] / v[o[1]], o[1],
                                                                          v[o[1]], pos[key]))
    ok(pos["avg"] in (4, 5), "R0: Gorsefield %d on the published average (4th or 5th)" % pos["avg"])
    v_nb = {c: fig[c]["nb"] for c in fig}
    ok(pos["nb"] >= 2 and (pos["nb"] > 2 or v_nb["QLBY"] / v_nb[PICK] >= 1.20),
       "R1: Gorsefield %d on the non-baseload rate, %.2f below Quillbury's" % (pos["nb"], v_nb["QLBY"] / v_nb[PICK]))
    ok(pos["pooled"] >= 3, "R2: Gorsefield %d on one marginal rate fitted over every hour" % pos["pooled"])
    ok(sum(1 for k in pos if pos[k] == 2) <= 1, "Gorsefield second on at most one rung: %s" % pos)
    ok(len({"BRKW", "QLBY", "TNFD", PICK}) == 4 and len(set(RUNG.values())) == 3, "each rung names its own store")
    pn = {c: fig[c]["pooled_ng"] for c in fig}
    ok(order(pn)[0] == "TNFD" and pn["TNFD"] / pn[order(pn)[1]] >= 1.15,
       "R2 with net generation as the regressor: Tennaford %.3f, %.3f over %s" % (pn["TNFD"], pn["TNFD"] / pn[order(pn)[1]],
                                                                                 order(pn)[1]))
    for key, want in (("hourly_average", "BRKW"), ("all_hours", "TNFD"), ("fossil", "TNFD")):
        v = {c: fig[c][key] for c in fig}
        o = order(v)
        ok(o[0] == want and o[0] != PICK, "reflex %s lands on %s (%.3f over %s %.3f), never Gorsefield" % (
            key, o[0], v[o[0]], o[1], v[o[1]]))
    # the grid: basis (average, non-baseload, marginal) by hours (all, the array's)
    cells = {"average, all hours": order({c: fig[c]["avg"] for c in fig})[0],
             "average, the array's hours": order({c: fig[c]["hourly_average"] for c in fig})[0],
             "non-baseload": order(v_nb)[0],
             "marginal, all hours": order({c: fig[c]["pooled"] for c in fig})[0],
             "marginal, the array's hours": go[0]}
    ok(cells == {"average, all hours": "BRKW", "average, the array's hours": "BRKW", "non-baseload": "QLBY",
                 "marginal, all hours": "TNFD", "marginal, the array's hours": PICK},
       "grid: only the margin in the array's hours reaches Gorsefield %s" % cells)
    ok(all(gold[PICK] >= 2.0 * gold[w] for w in RUNG.values()),
       "dominance: Gorsefield's %.3f is over twice every rung winner's own figure %s" % (
           gold[PICK], {w: round(gold[w], 3) for w in RUNG.values()}))
    ok(all(fig[w][k] - gold[w] >= 0.25 for k, w in RUNG.items()),
       "worth: every rung's winner avoids at least 0.25 t/MWh less than the rung says %s" % {
           w: (round(fig[w][k], 3), round(gold[w], 3)) for k, w in RUNG.items()})

    # ------------------------------------------------------------------ the estimator family (SF8)
    variants = {
        "month by hour": dict(bins="month_hour"), "hour of day": dict(bins="hour"), "UTC hour": dict(bins="utc_hour"),
        "demand": dict(regressor="demand"), "fossil generation": dict(regressor="fossil"),
        "month by hour, demand": dict(bins="month_hour", regressor="demand"),
        "hour of day, demand": dict(bins="hour", regressor="demand"),
        "month by hour, fossil": dict(bins="month_hour", regressor="fossil"),
        "hour of day, fossil": dict(bins="hour", regressor="fossil"),
        "imputed hours dropped": dict(drop_imputed=True),
        "register fuel rates": dict(rates_override={c: D["reg"][c] for c in D["reg"]}),
    }
    V = {lab: by_area(D, M.golden(D, **kw), "rate") for lab, kw in variants.items()}
    for c in W.ORDER:
        vals = [gold[c]] + [V[lab][c] for lab in V]
        ok(len({r2(v) for v in vals}) == 1 and min(edge(v, 0.01) for v in vals) >= 0.0005,
           "SF8: %s %.2f under all %d estimators (%.4f to %.4f, %.4f from a rounding edge)" % (
               c, gold[c], len(vals), min(vals), max(vals), min(edge(v, 0.01) for v in vals)))
    gaps = [gold[PICK] - gold[SECOND]] + [V[lab][PICK] - V[lab][SECOND] for lab in V]
    ok(len({r2(v) for v in gaps}) == 1 and min(edge(v, 0.01) for v in gaps) >= 0.0005,
       "SF8, SF12: the gap %.2f under every estimator (%.4f to %.4f), from unrounded rates" % (
           gaps[0], min(gaps), max(gaps)))
    ok(all(order(V[lab])[:2] == [PICK, SECOND] for lab in V), "SF8: every estimator picks Gorsefield, Halvane second")

    # ------------------------------------------------------------------ R1, the field's raw material ships
    hr = D["hourly"]
    ok(all(hr[c]["seen"].all() for c in W.ORDER) and rows[F["hourly"]] == 8 * 8760,
       "R1: the hourly operating data ship every area's net generation by fuel, demand and interchange for each of "
       "2025's 8,760 hours (%d rows)" % rows[F["hourly"]])
    heads = []
    for f in files:
        p = os.path.join(target, f)
        if f.endswith(".csv"):
            heads += open(p, encoding="utf-8").readline().strip().split(",")
        elif f.endswith(".xlsx"):
            from openpyxl import load_workbook
            wb = load_workbook(p, read_only=True)
            for ws in wb.worksheets:
                for r in ws.iter_rows(min_row=1, max_row=3, values_only=True):
                    heads += [str(c) for c in r if c is not None]
    ok(not any(re.search(r"marginal|avoided|margin\b|displaced", h, re.I) for h in heads),
       "R1: no shipped column or heading holds a marginal or avoided rate (%d headings read)" % len(heads))
    ok(rows[F["arrays"]] == 8 * 8760 and all(len(D["arrays"][st]) == 8760 for st in D["arrays"]),
       "R1: each array's modelled output ships hour by hour, so the weighting is built, never read")

    # ------------------------------------------------------------------ R2, each Layer 1 item mishandled
    def crit(fg):
        o = order(fg)
        return (o[0], o[1], r2(fg[PICK]), r2(fg[SECOND]), r2(fg[PICK] - fg[SECOND]), r2(fg["QLBY"]))
    base = crit(gold)
    for sh, what in ((-1, "the operating data's hour end read as the model's hour start"),
                     (1, "the model's hour start read as the operating data's hour end")):
        alt = by_area(D, M.golden(D, shift=sh), "rate")
        moved = {c: round(alt[c] - gold[c], 3) for c in W.ORDER if r2(alt[c]) != r2(gold[c])}
        ok(crit(alt) == base, "R2 clock: %s leaves the call, the first left out, the gap and components 1 to 3 "
                              "(%s); other stores move %s" % (what, base, moved))
    cc = by_area(D, M.golden_current_code_only(D))
    ok(crit(cc) == base and r2(cc["KSTP"]) == r2(gold["KSTP"]),
       "R2 renamed area: KPLP's rows dropped leave the call and components; Kestrel Point %.4f against %.4f" % (
           cc["KSTP"], gold["KSTP"]))
    ok(all(r2(V["imputed hours dropped"][c]) == r2(gold[c]) for c in W.ORDER),
       "R2 imputed demand: the operator's filled hours dropped move no figure at two decimals")
    D2 = copy.deepcopy(D)
    for c in W.ORDER:
        neg = np.clip(D2["hourly"][c]["solar"], None, 0.0)
        D2["hourly"][c]["solar"] = D2["hourly"][c]["solar"] - neg
        D2["hourly"][c]["ng"] = D2["hourly"][c]["ng"] - neg
    sc = by_area(D2, M.golden(D2), "rate")
    ok(all(r2(sc[c]) == r2(gold[c]) for c in W.ORDER) and
       all(np.allclose(M.emissions(D2["hourly"][c], D["rates"][c]["fuel"]), M.emissions(D["hourly"][c], D["rates"][c]["fuel"]))
           for c in W.ORDER),
       "R2 station service: solar's night readings clipped to zero change no CO2 and no figure")
    D3 = copy.deepcopy(D)
    for c in W.ORDER:
        D3["rates"][c]["fuel"] = {f: (v * M.LB_PER_T if v else v) for f, v in D["rates"][c]["fuel"].items()}
    lb = by_area(D3, M.golden(D3), "rate")
    ok(order(lb) == go and all(abs(lb[c] / gold[c] - M.LB_PER_T) < 1e-6 for c in W.ORDER),
       "R2 units: pounds per MWh left unconverted scale every figure alike; the order of the stores is unchanged")
    tmp = copy_without(target, {F["invoices"], F["rates24"], F["label"]})
    try:
        Dx = M.load(tmp)
        gx = by_area(Dx, M.golden(Dx), "rate")
        Rx = by_area(Dx, M.rungs(Dx))
    finally:
        shutil.rmtree(tmp)
    ok(gx == gold and all(Rx[c]["avg"] == fig[c]["avg"] and Rx[c]["nb"] == fig[c]["nb"] for c in W.ORDER),
       "R2 invoices, distractors: the call and every rung computed with the invoice export, the 2024 edition and the "
       "supplier's label deleted are identical (the main call reads none of them)")

    # ------------------------------------------------------------------ R3a, R3b on the clean truth
    if x is not None:
        T = truth_routes(x)
        tr = lambda k: order({c: T[c][k] for c in T})
        ok(tr("avg")[0] == "BRKW", "R3a: the easy route on the clean truth, the published average output emission rate "
                                   "taken as what each array avoids per MWh of its output in metric tonnes, files "
                                   "Brackwell (%.3f)" % T["BRKW"]["avg"])
        ok(tr("nb")[0] == "QLBY", "R3a: the published non-baseload rate on the clean truth files Quillbury (%.3f)" %
           T["QLBY"]["nb"])
        ok(tr("pooled")[0] == "TNFD", "R3a: one marginal rate per area fitted over every 2025 hour on the clean truth "
                                      "files Tennaford (%.3f)" % T["TNFD"]["pooled"])
        ok(tr("hourly_average")[0] == "BRKW" and tr("all_hours")[0] == "TNFD",
           "R3b: each lower rung's single correction stays wrong on the clean truth (the average in the array's hours "
           "Brackwell, the margin over all hours Tennaford)")
        ok(tr("decisive")[:2] == [PICK, SECOND],
           "R3b: the decisive step alone (the margin by season and hour, weighted by the array's output) files "
           "Gorsefield %.3f, Halvane %.3f" % (T[PICK]["decisive"], T[SECOND]["decisive"]))
        ok(tr("true")[:2] == [PICK, SECOND] and all(abs(T[c]["decisive"] - T[c]["true"]) < 0.01 for c in (PICK, SECOND)),
           "R3b: the generators the dispatch actually moved in each hour, weighted by the array's output, give "
           "Gorsefield %.3f and Halvane %.3f; the estimate is within 0.01 of both" % (T[PICK]["true"], T[SECOND]["true"]))
        ok(order({c: T[c]["true_all"] for c in T})[0] == "TNFD",
           "R3b: the true margin averaged over every hour is Tennaford's: the time of day, not the estimator, moves it")
        ok(all(order({c: T[c][k] for c in T})[0] == order({c: fig[c][k] for c in fig})[0]
               for k in ("avg", "nb", "pooled", "hourly_average")) and tr("decisive")[0] == go[0],
           "R3: every route files the same store on the shipped files as on the clean truth (the mess carries nothing)")

    # ------------------------------------------------------------------ R3c, nothing planted
    rep = []
    for c in W.ORDER:
        for f in ("coal", "natural gas", "petroleum", "nuclear", "hydro", "wind", "solar", "demand"):
            a = hr[c][f]
            if not np.any(a):
                continue
            days = a.reshape(365, 24)
            seen = {}
            for k, d in enumerate(days):
                if np.count_nonzero(d) < 6:
                    continue
                if (np.diff(d) == 0).all():
                    rep.append((c, f, k, "flat"))
                key = d.tobytes()
                if key in seen:
                    rep.append((c, f, k, "copy of day %d" % seen[key]))
                seen[key] = k
    for st, a in D["arrays"].items():
        days = a.reshape(365, 24)
        seen = {}
        for k, d in enumerate(days):
            if np.count_nonzero(d) >= 6:
                if d.tobytes() in seen:
                    rep.append((st, "array", k, "copy"))
                seen[d.tobytes()] = k
    ok(not rep, "R3c: no flat or copied day in any area's hourly series or any array's output %s" % rep[:5])
    docs = {"charter": text_of(os.path.join(target, F["charter"])), "label": text_of(os.path.join(target, F["label"])),
            "rates": text_of(os.path.join(target, F["rates"])), "rates24": text_of(os.path.join(target, F["rates24"])),
            "dictionary": text_of(os.path.join(target, F["dict"])), "manifest": text_of(os.path.join(target, F["manifest"]))}
    hits = METHOD_WORDS.findall(docs["charter"])
    ok(not hits, "R3c: the charter defines what is avoided and names no rate, basis or estimator %s" % hits)
    hits = AVOID_WORDS.findall(docs["rates"] + " " + docs["rates24"])
    ok(not hits, "R3c: neither edition of the publication says what a project avoids or names a margin %s" % hits)
    hits = DICT_WORDS.findall(docs["dictionary"] + " " + docs["manifest"])
    ok(not hits, "R3c: the dictionary and the manifest name no estimator, margin or avoided rate %s" % hits)
    ok(not AVOID_WORDS.findall(docs["label"]) and len(re.findall(r"pounds per megawatt-hour", docs["label"])) == 1,
       "R3c: the supplier's label carries one intensity for every store, ranks none and says nothing of avoiding")
    ok(sorted(files) == sorted(F.values()) and not any(f.endswith((".eml", ".msg", ".docx", ".txt", ".md")) for f in files),
       "R3c: no memo, email, draft or prior analysis ships; the folder is the declared %d files" % len(F))
    pub = M.read_rates(target)
    reg = M.register_rates(target)
    dif = max(abs(pub[c]["fuel"][f] - reg[c][f]) * M.LB_PER_T for c in reg for f in reg[c] if reg[c][f])
    ok(dif < 0.06, "R3c: the publication's fuel rates recompute from the register to %.3f lb/MWh (an honest product, "
                   "correct for what it defines)" % dif)
    plants = list(csv.DictReader(open(os.path.join(target, F["plants"]), encoding="utf-8")))
    worst = 0.0
    for c in W.ORDER:
        rs = [p for p in plants if p["balancing_area"] == c]
        gen = sum(float(p["net_generation_2025_mwh"]) for p in rs)
        co2 = sum(float(p["co2_2025_short_tons"]) for p in rs) * M.T_PER_SHORT_TON
        worst = max(worst, abs(co2 / gen - pub[c]["avg"]) * M.LB_PER_T)
    ok(worst < 0.06, "R3c: every area's published average output rate recomputes from the register (%.3f lb/MWh)" % worst)

    # ------------------------------------------------------------------ semantic forks
    for sf, doc, q in QUOTES:
        ok(q in docs[doc], "%s: ruling sentence in the %s: %r" % (sf, doc, q[:90]))
    share_out, same7 = [], True
    for st, s_ in D["stores"].items():
        off = D["codes"][s_["area"]]["utc"]
        a = D["arrays"][st]
        share_out.append(float(a[-abs(off):].sum() / a.sum()))
        local_year = float((g[st]["mer"] * g[st]["w"]).sum() / a.sum())        # the whole local year's output below
        same7 &= r2(local_year) == r2(g[st]["rate"]) and edge(local_year, 0.01) >= 0.0005
    ok(max(share_out) < 0.0001 and same7,
       "SF7: the array output in the local hours past the UTC year is at most %.4f%% of a store's 2025 output, so the "
       "local or the UTC year gives the same eight figures" % (100 * max(share_out)))
    # SF9's gross form: the model's local standard hour joined to the operating data's UTC label, the offset ignored.
    # Not a reading: every operating-data row states its local time and offset beside UTC, and the dictionary names
    # both clocks. It moves the call (the margin is a matter of the hour), so the files must close it.
    gross = {}
    for st, s_ in D["stores"].items():
        code = s_["area"]
        utc = D["codes"][code]["utc"]
        w = M.weights(D["arrays"][st], utc, shift=utc)
        gross[code] = float((g[st]["mer"] * w).sum() / w.sum())
    rows_h = []
    with open(os.path.join(target, F["hourly"]), newline="", encoding="utf-8") as fh:
        for k, r in enumerate(csv.DictReader(fh)):
            if k % 997 == 0:
                rows_h.append(r)
    off_ok = all(dt.datetime.fromisoformat(r["local_time_at_end_of_hour"]).replace(tzinfo=None) ==
                 dt.datetime.strptime(r["utc_time_at_end_of_hour"], "%Y-%m-%dT%H:%M:%SZ") +
                 dt.timedelta(hours=int(r["local_time_at_end_of_hour"][-6:-3])) for r in rows_h)
    ok(off_ok and "End of the hour on the area's local clock, with its UTC offset." in docs["dictionary"] and
       order(gross)[0] != PICK,
       "SF9 closed: every operating-data row states its local time with its offset (%d sampled) and the dictionary "
       "names both clocks, so the offset ignored (which files %s) is an error the files rule out, not a reading" % (
           len(rows_h), order(gross)[0]))
    codes = list(csv.DictReader(open(os.path.join(target, F["codes"]), encoding="utf-8")))
    ok([(r["ba_code"], r["former_code"], r["code_effective"]) for r in codes if r["former_code"]] ==
       [("KSTP", "KPLP", "2025-07-01")], "SF10: the code list files KPLP as Kestrel Point's former code until 1 July 2025")
    ok(edge(gold[PICK] - gold[SECOND], 0.01) >= 0.001 and abs(r2(gold[PICK]) - r2(gold[SECOND]) - r2(gold[PICK] - gold[SECOND])) < 1e-9,
       "SF12: the gap %.4f is the same from rounded and unrounded rates" % (gold[PICK] - gold[SECOND]))

    # ------------------------------------------------------------------ the asks
    Aa = by_area(D, M.asks(target, D, g))
    WP = by_area(D, M.ask_wrong_paths(target, D, g))
    for c in W.ORDER:
        a = Aa[c]
        ok(edge(a["use_mwh"], 1.0) >= 0.15 and edge(a["share"], 0.1) >= 0.03 and edge(a["price"], 1.0) >= 0.15,
           "asks %s: use %.3f MWh, share %.3f%%, price $%.3f, each clear of its rounding edge" % (
               c, a["use_mwh"], a["share"], a["price"]))
    use = {c: round(Aa[c]["use_mwh"]) for c in W.ORDER}
    for key, lab in (("kwh_all", "every kWh line summed"), ("max_line", "the largest kWh line per invoice"),
                     ("supplier", "the supplier's quantities alone"), ("both", "both vendors' energy lines")):
        miss = [c for c in W.ORDER if round(WP[c][key]) != use[c]]
        ok(len(miss) == 8, "ask A cracker path, %s: misses all eight stores (Brackwell %s against %s)" % (
            lab, "{:,}".format(round(WP["BRKW"][key])), "{:,}".format(use["BRKW"])))
    hi = {c: 100.0 * (WP[c]["supplier"] / Aa[c]["use_mwh"] - 1.0) for c in W.ORDER}
    ok(all(3.5 <= v <= 6.5 for v in hi.values()), "ask A hazard H1: the supplier's quantities run %.1f to %.1f per cent "
                                                    "above the meters" % (min(hi.values()), max(hi.values())))
    rb = [c for c in W.ORDER if round(WP[c]["all_bills"]) != use[c]]
    ok(rb == ["QLBY", "OSTM"], "ask A hazard H2: keeping the estimated bills moves Quillbury and Ostlemere only %s" % rb)
    idt = [c for c in W.ORDER if round(WP[c]["by_invoice_date"]) != use[c]]
    ok(len(idt) == 8, "ask A hazard H4: counting invoices dated in 2025 misses %d stores" % len(idt))
    rows_inv = M.read_invoices(target)
    keep = M.current_bills(rows_inv)
    by_inv = {}
    for r in rows_inv:
        by_inv.setdefault(r["invoice_number"], []).append(r)
    per_meter = {}
    for no, L in by_inv.items():
        if no not in keep or L[0]["vendor_name"] == W.SUPPLIER:
            continue
        frac = M.in_2025(L[0]["d0"], L[0]["d1"])
        for m_, q in M.metered_kwh(L).items():
            k = (L[0]["store"], m_)
            per_meter[k] = per_meter.get(k, 0.0) + q * frac / 1000.0
    one = {}
    for st, s_ in D["stores"].items():
        ms = [v for (st_, m_), v in per_meter.items() if st_ == st]
        one[s_["area"]] = round(max(ms))
    h3 = [c for c in W.ORDER if one[c] != use[c]]
    ok(h3 == ["QLBY", "KSTP"], "ask A hazard H3: one meter per store moves Quillbury and Kestrel Point only %s" % h3)
    price = {c: round(Aa[c]["price"]) for c in W.ORDER}
    pg = round(WP["HLVN"]["paid_with_gas"] / Aa["HLVN"]["use_mwh"])
    ok(pg != price["HLVN"] and all(round(WP[c]["paid_with_gas"] / Aa[c]["use_mwh"]) == price[c] for c in W.ORDER if c != "HLVN"),
       "ask C device: the gas lines on Halvane's invoices move only Halvane's price ($%d against $%d)" % (pg, price["HLVN"]))
    pu = [c for c in W.ORDER if round(WP[c]["paid_util_only"] / Aa[c]["use_mwh"]) != price[c]]
    ok(len(pu) == 8, "ask C hazard: the utility's dollars alone miss all eight prices")
    D0 = copy.deepcopy(D)
    for c in W.ORDER:
        for k in ("demand", "ng", "coal", "natural gas", "petroleum"):
            D0["hourly"][c][k] = np.zeros(8760)
    A0 = by_area(D0, M.asks(target, D0, g))
    ok(all(A0[c]["use_mwh"] == Aa[c]["use_mwh"] and A0[c]["price"] == Aa[c]["price"] for c in W.ORDER),
       "ask separation: the asks read no operating-data row (zeroed, every ask figure is unchanged)")

    # ------------------------------------------------------------------ realism sweep
    for key, lab in (("use_mwh", "use"), ("price", "price"), ("share", "share")):
        vals = [round(Aa[c][key], 1) for c in W.ORDER]
        ok(len(set(vals)) == 8, "realism: no two stores share a %s %s" % (lab, vals))
    ok(not any(round(Aa[c]["use_mwh"]) % 100 == 0 for c in W.ORDER), "realism: no store's use is a round hundred MWh")
    cnt = [sum(1 for r in rows_inv if r["store"] == st) for st in D["stores"]]
    ok(len(set(cnt)) >= 6, "realism: invoice line counts differ across stores %s" % cnt)
    raw = b"".join(open(os.path.join(target, f), "rb").read() for f in files if not f.endswith(".csv"))
    tells = [t for t in (b"ReportLab", b"reportlab", b"openpyxl", b"python", b"Python") if t in raw]
    ok(not tells, "realism: no writing library's name in any document or workbook %s" % tells)
    stamps = {os.path.getmtime(os.path.join(target, f)) for f in files}
    ok(len(stamps) == 1, "realism: every file carries the folder's one export time")
    idates = [r["idate"] for r in rows_inv]
    ok(max(idates) <= W.AP_PULL, "realism: no invoice dated after the accounts-payable pull (%s)" % max(idates))
    names = set(re.findall(r"\b(Barbara Rangel|Russell Scott)\b", docs["charter"] + " " + docs["dictionary"]))
    ok(names == {"Barbara Rangel"} and "Russell Scott" not in docs["charter"],
       "realism: the charter's owner is the drawn persona (Barbara Rangel); no other person is named in the documents")

    # ------------------------------------------------------------------ the goldens
    memo = os.path.join(golden_dir, W.GOLDEN["memo"])
    hourly = os.path.join(golden_dir, W.GOLDEN["hourly"])
    if os.path.exists(memo):
        mt = text_of(memo)
        ok(("Site the 2027 solar pilot at Corvane Gorsefield" in mt) and ("%.2f" % gold[PICK]) in mt,
           "golden memo: the title names Gorsefield and its %.2f" % gold[PICK])
        ok(all(("%.2f" % gold[c]) in mt for c in W.ORDER), "golden memo: every store's figure is in the table")
        ok(not METHOD_WORDS.findall(open(os.path.join(TASK, "prompt.txt")).read()),
           "prompt: names no basis, rate or estimator")
    if os.path.exists(hourly):
        acc = {}
        n = 0
        with open(hourly, newline="", encoding="utf-8") as f:
            for r in csv.DictReader(f):
                n += 1
                a = acc.setdefault(r["balancing_area"], [0.0, 0.0])
                a[0] += float(r["avoided_co2_kg"])
                a[1] += float(r["array_output_kwh"])
        ok(n == 70080, "golden hourly file: %d rows, one per store and local standard hour" % n)
        ok(all(r2(acc[c][0] / acc[c][1]) == r2(gold[c]) for c in W.ORDER),
           "golden hourly file: each store's avoided CO2 over its output equals its memo figure")
    if fig_path and os.path.exists(fig_path):
        Fj = json.load(open(fig_path))
        ok(Fj["call"] == "Gorsefield" and Fj["first_left_out"] == "Halvane" and abs(Fj["gap"] - r2(gold[PICK] - gold[SECOND])) < 1e-9,
           "figures.json: Gorsefield, Halvane and the gap agree with the checks")

    print("%d checks, %d failed" % (COUNT[0], len(FAILS)))
    return not FAILS


if __name__ == "__main__":
    sys.path.insert(0, HERE)
    import build as BD
    x = BD.world()
    sys.exit(0 if run(os.path.join(TASK, "target"), os.path.join(TASK, "golden"), x=x,
                      meta_path=os.path.join(TASK, "metadata.json"),
                      fig_path=os.path.join(HERE, "figures.json")) else 1)
