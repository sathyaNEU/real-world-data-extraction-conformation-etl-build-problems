"""Everything the build computes from the shipped files: the golden (each area's marginal emission rate by season
and local standard hour, from the hour-to-hour change in its fossil CO2 against the change in its net generation,
weighted by each store's modelled array output), every rung and reflex of the ladder, the estimator variants that
must agree with the golden, the natural mishandlings that must leave the call where it is, and the asks from the
invoices."""
import csv
import datetime as dt
import os

import numpy as np

import world as W

H = 8760
LB_PER_T = 2204.62
T_PER_SHORT_TON = 0.90718474
UTC0 = dt.datetime(2025, 1, 1)
SEASON = {12: 0, 1: 0, 2: 0, 3: 1, 4: 1, 5: 1, 6: 2, 7: 2, 8: 2, 9: 3, 10: 3, 11: 3}
FUELS = ("coal", "natural gas", "petroleum")
COLS = {"coal": "ng_coal_mw", "natural gas": "ng_natural_gas_mw", "petroleum": "ng_petroleum_mw"}


def _rows(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def _num(x):
    return float(x) if x not in ("", None) else 0.0


# ---------------------------------------------------------------------------------------------- reading the pack
def read_codes(target):
    codes = {}
    for r in _rows(os.path.join(target, W.FILES["codes"])):
        codes[r["ba_code"]] = dict(utc=int(r["utc_offset"][:3]), former=r["former_code"] or None)
    return codes


def read_hourly(target, codes):
    """Per area (current code): arrays over the 8,760 UTC hours of 2025, indexed by the hour's start."""
    alias = {c["former"]: k for k, c in codes.items() if c["former"]}
    out = {k: {f: np.zeros(H) for f in ("demand", "ng", "ti", "coal", "natural gas", "petroleum", "nuclear", "hydro",
                                        "wind", "solar")} for k in codes}
    for k in codes:
        out[k]["imputed"] = np.zeros(H, bool)
        out[k]["seen"] = np.zeros(H, bool)
        out[k]["former"] = np.zeros(H, bool)
    for r in _rows(os.path.join(target, W.FILES["hourly"])):
        code = alias.get(r["balancing_area"], r["balancing_area"])
        end = dt.datetime.strptime(r["utc_time_at_end_of_hour"], "%Y-%m-%dT%H:%M:%SZ")
        h = int((end - UTC0).total_seconds() // 3600) - 1
        a = out[code]
        a["demand"][h] = _num(r["demand_mw"])
        a["ng"][h] = _num(r["net_generation_mw"])
        a["ti"][h] = _num(r["total_interchange_mw"])
        for f, col in (("coal", "ng_coal_mw"), ("natural gas", "ng_natural_gas_mw"), ("petroleum", "ng_petroleum_mw"),
                       ("nuclear", "ng_nuclear_mw"), ("hydro", "ng_hydro_mw"), ("wind", "ng_wind_mw"),
                       ("solar", "ng_solar_mw")):
            a[f][h] = _num(r[col])
        a["imputed"][h] = r["demand_imputed"] == "Y"
        a["seen"][h] = True
        a["former"][h] = r["balancing_area"] != code
    for k in codes:
        assert out[k]["seen"].all(), k
    return out


def read_rates(target, name="rates"):
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(target, W.FILES[name]), data_only=True)
    ws = wb.worksheets[0]
    hdr = [c.value for c in ws[3]]
    out = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row[0]:
            continue
        d = dict(zip(hdr, row))
        t = lambda v: (v / LB_PER_T) if v is not None else None
        out[d["BA code"]] = dict(avg=t(d["CO2 output emission rate (lb/MWh)"]),
                                 fossil=t(d["CO2 fossil output emission rate (lb/MWh)"]),
                                 nb=t(d["CO2 non-baseload output emission rate (lb/MWh)"]),
                                 fuel={"coal": t(d["CO2 coal output emission rate (lb/MWh)"]),
                                       "natural gas": t(d["CO2 gas output emission rate (lb/MWh)"]),
                                       "petroleum": t(d["CO2 oil output emission rate (lb/MWh)"])})
    return out


def register_rates(target):
    """Fuel output rates recomputed from the plant register (short tons to tonnes)."""
    acc = {}
    for r in _rows(os.path.join(target, W.FILES["plants"])):
        f = r["primary_fuel"]
        if f not in FUELS:
            continue
        k = acc.setdefault(r["balancing_area"], {}).setdefault(f, [0.0, 0.0])
        k[0] += _num(r["net_generation_2025_mwh"])
        k[1] += _num(r["co2_2025_short_tons"]) * T_PER_SHORT_TON
    return {a: {f: (v[1] / v[0] if v[0] > 0 else None) for f, v in d.items()} for a, d in acc.items()}


def read_stores(target):
    umap = {r["utility_name"]: r["balancing_area"] for r in _rows(os.path.join(target, W.FILES["utilmap"]))}
    out = {}
    for r in _rows(os.path.join(target, W.FILES["sites"])):
        out[int(r["store_number"])] = dict(name=r["store_name"], town=r["city"], area=umap[r["delivery_utility"]],
                                           dc=float(r["array_kw_dc"]))
    return out


def read_arrays(target):
    out = {}
    for r in _rows(os.path.join(target, W.FILES["arrays"])):
        out.setdefault(int(r["store_number"]), []).append(float(r["ac_output_kw"]))
    return {k: np.array(v) for k, v in out.items()}


def load(target):
    codes = read_codes(target)
    return dict(codes=codes, hourly=read_hourly(target, codes), rates=read_rates(target), reg=register_rates(target),
                stores=read_stores(target), arrays=read_arrays(target))


# ---------------------------------------------------------------------------------------------- clocks
def clock(utc):
    """For an area: each UTC hour's local standard hour (beginning), month and day of the local standard date."""
    t = [UTC0 + dt.timedelta(hours=h + utc) for h in range(H)]
    return np.array([x.hour for x in t]), np.array([x.month for x in t]), np.array([x.timetuple().tm_yday for x in t])


def weights(arr, utc, shift=0):
    """The store's array output placed on the area's UTC hours: local standard hour k is UTC hour k - utc. shift moves
    every hour (an hour's end read as its beginning, or the reverse)."""
    w = np.zeros(H)
    k = np.arange(H)
    h = k - utc + shift
    ok = (h >= 0) & (h < H)
    np.add.at(w, h[ok], arr[ok])
    return w


# ---------------------------------------------------------------------------------------------- the estimators
def emissions(a, rates):
    return sum(a[f] * (rates.get(f) or 0.0) for f in FUELS)


def binned_slope(dy, dx, key, mask=None):
    """OLS slope of dy on dx with an intercept, within each bin of key; every hour gets its bin's slope."""
    out = np.zeros(len(dy))
    m_all = np.ones(len(dy), bool) if mask is None else mask
    for k in np.unique(key):
        m = (key == k)
        mm = m & m_all
        x, y = dx[mm] - dx[mm].mean(), dy[mm] - dy[mm].mean()
        v = (x * x).sum()
        out[m] = (x * y).sum() / v if v > 0 else 0.0
    return out


def mer(a, utc, rates, bins="season_hour", regressor="ng", drop_imputed=False):
    """The marginal emission rate for every UTC hour of 2025: the slope of the hour-to-hour change in the area's fossil
    CO2 on the change in its net generation (or demand, or fossil generation), within bins of the local hour."""
    E = emissions(a, rates)
    lh, mon, _ = clock(utc)
    if bins == "utc_hour":
        lh = (lh - utc) % 24
    x = {"demand": a["demand"], "fossil": sum(a[f] for f in FUELS), "ng": a["ng"]}[regressor]
    dE = np.diff(E, prepend=E[0])
    dX = np.diff(x, prepend=x[0])
    season = np.array([SEASON[m] for m in mon])
    key = {"season_hour": season * 24 + lh, "month_hour": mon * 24 + lh, "hour": lh, "utc_hour": season * 24 + lh,
           "all": np.zeros(H, int)}[bins]
    mask = np.ones(H, bool)
    mask[0] = False
    if drop_imputed:
        imp = a["imputed"] | np.roll(a["imputed"], 1)
        mask &= ~imp
    return binned_slope(dE, dX, key, mask)


def golden(D, **kw):
    """Per store: the avoided CO2 per MWh of its array's 2025 output, the output and the tonnes avoided."""
    out = {}
    for st, s in D["stores"].items():
        code = s["area"]
        utc = D["codes"][code]["utc"]
        m = mer(D["hourly"][code], utc, kw.get("rates_override", {}).get(code, D["rates"][code]["fuel"]),
                bins=kw.get("bins", "season_hour"), regressor=kw.get("regressor", "ng"),
                drop_imputed=kw.get("drop_imputed", False))
        w = weights(D["arrays"][st], utc, kw.get("shift", 0))
        out[st] = dict(area=code, rate=float((m * w).sum() / w.sum()), mwh=float(D["arrays"][st].sum() / 1000.0),
                       tonnes=float((m * w).sum() / 1000.0), mer=m, w=w)
    return out


def golden_current_code_only(D):
    """As golden, with the former-code months of the renamed area out of every regression bin."""
    out = {}
    for st, s in D["stores"].items():
        code = s["area"]
        a = D["hourly"][code]
        utc = D["codes"][code]["utc"]
        E = emissions(a, D["rates"][code]["fuel"])
        lh, mon, _ = clock(utc)
        dE = np.diff(E, prepend=E[0])
        dX = np.diff(a["ng"], prepend=a["ng"][0])
        season = np.array([SEASON[m] for m in mon])
        mask = ~a["former"] & ~np.roll(a["former"], 1)
        mask[0] = False
        m = binned_slope(dE, dX, season * 24 + lh, mask)
        w = weights(D["arrays"][st], utc)
        if a["former"].any():
            w = np.where(a["former"], 0.0, w)
        out[st] = float((m * w).sum() / w.sum())
    return out


# ---------------------------------------------------------------------------------------------- the rungs
def rungs(D):
    """Each easy route's figure per store."""
    g = golden(D)
    out = {}
    for st, s in D["stores"].items():
        code = s["area"]
        a = D["hourly"][code]
        utc = D["codes"][code]["utc"]
        r = D["rates"][code]
        E = emissions(a, r["fuel"])
        w = g[st]["w"]
        dE = np.diff(E, prepend=E[0])
        dD = np.diff(a["demand"], prepend=a["demand"][0])
        x, y = dD[1:] - dD[1:].mean(), dE[1:] - dE[1:].mean()
        pooled = float((x * y).sum() / (x * x).sum())
        dN = np.diff(a["ng"], prepend=a["ng"][0])
        xn = dN[1:] - dN[1:].mean()
        pooled_ng = float((xn * y).sum() / (xn * xn).sum())
        intensity = np.where(a["ng"] > 0, E / a["ng"], 0.0)
        out[st] = dict(area=code, avg=r["avg"], nb=r["nb"], fossil=r["fossil"], pooled=pooled, pooled_ng=pooled_ng,
                       all_hours=float(g[st]["mer"].mean()), hourly_average=float((intensity * w).sum() / w.sum()),
                       golden=g[st]["rate"])
    return out


def ranking(fig, key):
    return sorted(fig, key=lambda st: -fig[st][key])


# ---------------------------------------------------------------------------------------------- the asks
def read_invoices(target):
    rows = _rows(os.path.join(target, W.FILES["invoices"]))
    for r in rows:
        r["qty"] = float(r["quantity"]) if r["quantity"] else None
        r["amt"] = float(r["line_amount_usd"])
        r["d0"] = dt.date.fromisoformat(r["service_period_start"])
        r["d1"] = dt.date.fromisoformat(r["service_period_end"])
        r["idate"] = dt.date.fromisoformat(r["invoice_date"])
        r["store"] = int(r["store_number"])
    return rows


def in_2025(d0, d1):
    a, b = max(d0, dt.date(2025, 1, 1)), min(d1, dt.date(2025, 12, 31))
    return max((b - a).days + 1, 0) / ((d1 - d0).days + 1)


def current_bills(rows):
    """The invoices that stand: for each vendor, meter and service period, the latest invoice (an invoice billed again on
    an actual read replaces the estimated one)."""
    inv = {}
    for r in rows:
        inv.setdefault(r["invoice_number"], r)
    latest = {}
    for no, r in inv.items():
        key = (r["vendor_number"], r["meter_number"], r["d0"], r["d1"])
        if key not in latest or r["idate"] > inv[latest[key]]["idate"]:
            latest[key] = no
    return set(latest.values())


DIST_ENERGY = ("Distribution energy charge", "Delivery service - energy", "Distribution kWh charge",
               "Energy delivery charge", "Distribution charge per kWh", "Delivery energy charge", "Distribution energy",
               "Electric distribution energy charge")


def metered_kwh(lines):
    """An invoice's metered kWh per meter: its distribution energy quantity (summed over the declining blocks where the
    utility bills distribution energy in blocks)."""
    out = {}
    for r in lines:
        d = r["line_description"]
        if r["unit_of_measure"] == "KWH" and (d in DIST_ENERGY or d.startswith("Distribution kWh charge,")):
            out[r["meter_number"]] = out.get(r["meter_number"], 0.0) + r["qty"]
    return out


def asks(target, D, golden_fig):
    rows = read_invoices(target)
    keep = current_bills(rows)
    by_inv = {}
    for r in rows:
        by_inv.setdefault(r["invoice_number"], []).append(r)
    out = {}
    for st, s in D["stores"].items():
        use = 0.0
        months = np.zeros(12)
        paid = 0.0
        for no, L in by_inv.items():
            if no not in keep or L[0]["store"] != st:
                continue
            frac = in_2025(L[0]["d0"], L[0]["d1"])
            paid += frac * sum(r["amt"] for r in L if not r["line_description"].startswith("Gas"))
            if L[0]["vendor_name"] == W.SUPPLIER:
                continue
            q = sum(metered_kwh(L).values())
            use += q * frac
            d0, d1 = L[0]["d0"], L[0]["d1"]
            n = (d1 - d0).days + 1
            for k in range(n):
                day = d0 + dt.timedelta(days=k)
                if day.year == 2025:
                    months[day.month - 1] += q / n / 1000.0
        arr_m = np.zeros(12)
        t0 = dt.datetime(2025, 1, 1)
        arr = D["arrays"][st]
        for k in range(H):
            arr_m[(t0 + dt.timedelta(hours=k)).month - 1] += arr[k] / 1000.0
        out[st] = dict(use_mwh=use / 1000.0, paid=paid, price=paid / (use / 1000.0),
                       share=100.0 * golden_fig[st]["mwh"] / (use / 1000.0), months=months, array_months=arr_m)
    return out


def ask_wrong_paths(target, D, golden_fig):
    """The natural and half-repaired readings of the invoice asks, each a complete pipeline: every kWh line summed;
    one kWh line per invoice (the largest); the supplier's calendar months; both vendors' energy lines; every bill
    including the estimated ones billed again; invoices dated in 2025; dollars from the utility alone, or with the
    gas lines."""
    rows = read_invoices(target)
    keep = current_bills(rows)
    by_inv = {}
    for r in rows:
        by_inv.setdefault(r["invoice_number"], []).append(r)
    out = {}
    for st in D["stores"]:
        acc = dict(kwh_all=0.0, max_line=0.0, supplier=0.0, both=0.0, all_bills=0.0, by_invoice_date=0.0,
                   paid_util_only=0.0, paid_with_gas=0.0, paid_all_bills=0.0)
        for no, L in by_inv.items():
            if L[0]["store"] != st:
                continue
            frac = in_2025(L[0]["d0"], L[0]["d1"])
            util = L[0]["vendor_name"] != W.SUPPLIER
            kwh = [r["qty"] for r in L if r["unit_of_measure"] == "KWH"]
            q = sum(metered_kwh(L).values()) if util else 0.0
            if util:
                acc["all_bills"] += q * frac
            if no not in keep:
                acc["paid_all_bills"] += frac * sum(r["amt"] for r in L if not r["line_description"].startswith("Gas"))
                continue
            acc["kwh_all"] += frac * sum(kwh)
            acc["max_line"] += frac * (max(kwh) if kwh else 0.0)
            if util:
                acc["both"] += q * frac
                if L[0]["idate"].year == 2025:
                    acc["by_invoice_date"] += q
                acc["paid_util_only"] += frac * sum(r["amt"] for r in L if not r["line_description"].startswith("Gas"))
            else:
                e = sum(r["qty"] for r in L if r["line_description"].startswith("Energy"))
                acc["supplier"] += e * frac
                acc["both"] += e * frac
            acc["paid_with_gas"] += frac * sum(r["amt"] for r in L)
            acc["paid_all_bills"] += frac * sum(r["amt"] for r in L if not r["line_description"].startswith("Gas"))
        out[st] = {k: (v / 1000.0 if not k.startswith("paid") else v) for k, v in acc.items()}
    return out
