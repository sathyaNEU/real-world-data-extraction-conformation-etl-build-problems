"""The golden deliverables: the 2025-26 inventory workbook (the residential wood combustion inventory by burn zone and
SCC, the propane comparison it rests on, the commercial and institutional propane category, the county outside the
maintenance area on the state tool) and the zone determinations file; returns every figure the write-up states."""
import csv
import math

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import golden as G

BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")


def whole(x):
    return int(math.floor(x + 0.5))


def r1(x):
    return math.floor(x * 10 + 0.5) / 10.0


def r2(x):
    return math.floor(x * 100 + 0.5) / 100.0


def hundred(x):
    return int(math.floor(x / 100.0 + 0.5) * 100)


def _head(ws, row, wrap=True):
    for c in ws[row]:
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=wrap, vertical="top")


def write(path, p):
    homes = G.propane_homes(p)
    inv = G.inventory(p, homes=homes)
    dflt = G.inventory(p, "default")
    a = G.ask_a(p)
    b = {c: v for c, v in G.ask_b(p).items() if c != "_main"}
    zones = p["zones"]
    over = inv["over"]
    share = p["fund"] // len(over) if over else 0
    figs = dict(call=dict(over=over, fund=p["fund"], share=share), zones={}, steps={}, ask_a={}, ask_b={},
                order=list(zones), scc_codes=list(p["scc_of"].values()))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RWC inventory"
    ws.append(["Residential wood combustion, 2025-26 heating season: %s over their Table 4-2 allotments" %
               (", ".join(over[:-1]) + " and " + over[-1] if len(over) > 1 else (over[0] if over else "no zone"))])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Halvard Valley PM2.5 maintenance area, 1 October 2025 to 30 April 2026. Wood burned is the 2024-25 "
               "survey's per appliance plus the wood burned in place of propane (sheet Propane comparison); PM2.5 at "
               "the 2023 factors. Whole tons; zone totals from unrounded cells."])
    ws.append([])
    ws.append(["Burn zone", "SCC", "SCC description", "Appliances in use", "Wood burned (tons)",
               "of which in place of propane (tons)", "PM2.5 (tons)"])
    _head(ws, 4)
    for z in zones:
        for c, scc in p["scc_of"].items():
            x = inv["cells"][z, c]
            extra_t = G.tons(p, c, x["extra"])
            ws.append([z, scc, p["desc_of"][c], x["n"], whole(x["wood"]), whole(extra_t), whole(x["pm"])])
        ws.append([z + " total", None, None, sum(inv["cells"][z, c]["n"] for c in p["scc_of"]),
                   whole(sum(inv["cells"][z, c]["wood"] for c in p["scc_of"])),
                   whole(sum(G.tons(p, c, inv["cells"][z, c]["extra"]) for c in p["scc_of"])), whole(inv["zone"][z]["pm"])])
        for c in ws[ws.max_row]:
            c.font = BOLD
            c.fill = TOTAL_FILL
    ws.append(["Maintenance area", None, None, sum(v["n"] for v in inv["cells"].values()),
               whole(sum(v["wood"] for v in inv["cells"].values())),
               whole(sum(G.tons(p, k[1], v["extra"]) for k, v in inv["cells"].items())),
               whole(sum(inv["zone"][z]["pm"] for z in zones))])
    for c in ws[ws.max_row]:
        c.font = BOLD
    for i, w in enumerate([20, 12, 52, 12, 13, 16, 11], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=5):
        for c in row[3:]:
            c.number_format = "#,##0"
    ws.freeze_panes = "D5"
    ws.print_title_rows = "4:4"

    pc = wb.create_sheet("Propane comparison")
    pc.append(["Residential propane accounts in each zone, in-season deliveries (1 October to 30 April), gallons. "
               "Accounts matched to parcels and to the registry's heating appliances in use in both seasons."])
    pc["A1"].font = BOLD
    pc.append([])
    pc.append(["Burn zone", "Group", "Accounts", "2023-24", "2024-25", "2025-26", "2024-25 / 2023-24",
               "2025-26 / 2024-25", "Gallons not delivered", "Wood in their place (tons)"])
    _head(pc, 3)
    lo_c, hi_c, lo_t, hi_t, gap = 9.0, 0.0, 9.0, 0.0, 0.0
    for z in zones:
        comp = [h for h in homes if h["zone"] == z and h["cls"] is None]
        sums = [sum(h["g"][t] for h in comp) for t in (2023, 2024, 2025)]
        rc_pre, rc = sums[1] / sums[0], sums[2] / sums[1]
        pc.append([z, "No heating appliance", len(comp)] + [round(s) for s in sums] + [round(rc_pre, 4), round(rc, 4),
                                                                                       None, None])
        trt_all = [h for h in homes if h["zone"] == z and h["cls"]]
        st = [sum(h["g"][t] for h in trt_all) for t in (2023, 2024, 2025)]
        rt_pre, rt = st[1] / st[0], st[2] / st[1]
        lo_c, hi_c = min(lo_c, rc), max(hi_c, rc)
        lo_t, hi_t = min(lo_t, 1 - rt), max(hi_t, 1 - rt)
        gap = max(gap, abs(rt_pre - rc_pre))
        for c, scc in p["scc_of"].items():
            trt = [h for h in trt_all if h["cls"] == c]
            if not trt:
                continue
            s = [sum(h["g"][t] for h in trt) for t in (2023, 2024, 2025)]
            x = inv["cells"][z, c]
            pc.append([z, p["desc_of"][c], len(trt)] + [round(v) for v in s] + [round(s[1] / s[0], 4),
                                                                                 round(s[2] / s[1], 4), round(x["gal"]),
                                                                                 whole(G.tons(p, c, x["extra"]))])
        pc.append([z + ", heating appliance", "All", len(trt_all)] + [round(v) for v in st] +
                  [round(rt_pre, 4), round(rt, 4), round(sum(inv["cells"][z, c]["gal"] for c in p["scc_of"])),
                   whole(sum(G.tons(p, c, inv["cells"][z, c]["extra"]) for c in p["scc_of"]))])
        for cc in pc[pc.max_row]:
            cc.font = BOLD
            cc.fill = TOTAL_FILL
    for i, w in enumerate([24, 50, 10, 12, 12, 12, 11, 11, 13, 13], 1):
        pc.column_dimensions[get_column_letter(i)].width = w
    for row in pc.iter_rows(min_row=4):
        for c in row[2:]:
            if c.column in (7, 8):
                c.number_format = "0.0000"
            else:
                c.number_format = "#,##0"
    pc.freeze_panes = "C4"

    cp = wb.create_sheet("Commercial propane")
    cp.append(["Commercial and institutional propane delivered in the 2025-26 heating season, by delivery date"])
    cp["A1"].font = BOLD
    cp.append(["Wapiti, Coldwater and cooperative accounts classed commercial on parcels in the burn zones; NOx at the "
               "2023 factor of %.3f lb/MMBtu and %s Btu a gallon." % (p["ef_lp"]["Commercial/institutional propane "
                                                                       "combustion", "NOx"], format(int(p["btu_gal"]), ","))])
    cp.append(["Burn zone", "SCC", "Gallons delivered", "NOx (tons)"])
    _head(cp, 3)
    for z in zones:
        cp.append([z, "2103007000", hundred(a[z]["gal"]), r2(a[z]["nox"])])
        figs["ask_a"][z] = dict(gal=hundred(a[z]["gal"]), nox=r2(a[z]["nox"]))
    for i, w in enumerate([20, 12, 18, 12], 1):
        cp.column_dimensions[get_column_letter(i)].width = w
    for row in cp.iter_rows(min_row=4):
        row[2].number_format = "#,##0"
        row[3].number_format = "0.00"

    ot = wb.create_sheet("Outside the area")
    ot.append(["Halvard County outside the maintenance area, 2025-26, on the state tool (2023 edition)"])
    ot["A1"].font = BOLD
    ot.append(["%s main-heat households outside the area (ACS 2019-2023, B25040 wood), %s secondary households "
               "for each." % (format(whole(G.main_households(p)), ","), ("%.1f" % p["secondary"]))])
    ot.append(["SCC", "SCC description", "Appliances", "Wood burned (tons)", "PM2.5 (tons)"])
    _head(ot, 3)
    for c, scc in p["scc_of"].items():
        ot.append([scc, p["desc_of"][c], whole(b[c]["app"]), whole(b[c]["wood"]), r1(b[c]["pm"])])
        figs["ask_b"][scc] = dict(app=whole(b[c]["app"]), pm=r1(b[c]["pm"]))
    ot.append(["Total", None, whole(sum(v["app"] for v in b.values())), whole(sum(v["wood"] for v in b.values())),
               r1(sum(v["pm"] for v in b.values()))])
    for c in ot[ot.max_row]:
        c.font = BOLD
    for i, w in enumerate([12, 52, 12, 16, 12], 1):
        ot.column_dimensions[get_column_letter(i)].width = w
    for row in ot.iter_rows(min_row=4):
        row[2].number_format = "#,##0"
        row[3].number_format = "#,##0"
        row[4].number_format = "#,##0.0"

    nt = wb.create_sheet("Notes")
    notes = [
        "Prepared by Martin Booth, Halvard County Air Quality Program, 27 July 2026, for the 2026-27 changeout "
        "season's determinations (County Code 8.30.140).",
        "Appliances: the county registry's solid fuel devices in use in the season (installed before 1 October 2025 "
        "and not removed), by device code and its SCC.",
        "Wood per appliance: the 2024-25 household heating survey's estimate for the zone and class (methods note: "
        "weighted zone mean, classes under 10 respondents at the five-zone mean); face cords at a third of a cord, "
        "pickup loads at half, 40-pound bags, 1.163 tons a cord.",
        "Wood in place of propane: each zone's propane homes with a heating appliance, their 2024-25 gallons at the "
        "ratio the zone's propane homes without one moved by, less their 2025-26 gallons; the heat at 91,500 Btu a "
        "gallon and 85 per cent furnace efficiency, burned at the appliance's seasonal efficiency (18.5 million Btu a "
        "cord, 16.4 a ton of pellets). Fireplaces carry no heating efficiency in the state tool.",
        "Commercial propane: Coldwater statement accounts post a month's deliveries on the first of the next month; "
        "their delivery dates are the ticket reference's. NOx at 0.142 lb/MMBtu.",
        "Outside the area: the 2019-2023 ACS release the tool names; the assessor's 2010 block groups carried to the "
        "2020 block groups through the Census relationship file.",
        "Tons rounded for display (whole tons, the county's appliances whole and their PM2.5 to one decimal); "
        "determinations on unrounded values.",
    ]
    for line in notes:
        nt.append([line])
        nt.cell(row=nt.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    nt.column_dimensions["A"].width = 120
    wb.save(path("inventory"))

    with open(path("determinations"), "w", newline="") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(["burn_zone", "pm25_2025_26_tons", "allotment_tons", "exceeded_allotment", "fund_share_usd",
                    "changeout_season"])
        for z in zones:
            x = inv["zone"][z]
            sh = share if z in over else 0
            w.writerow([z, whole(x["pm"]), "%.1f" % x["allot"], "Yes" if x["over"] else "No", sh, "2026-27"])
            figs["zones"][z] = dict(pm=whole(x["pm"]), allot=x["allot"], exceeded=x["over"], share=sh,
                                    extra_pm=whole(x["pm"] - dflt["zone"][z]["pm"]), default_pm=whole(dflt["zone"][z]["pm"]))
    figs["constants"] = dict(btu_gal=int(p["btu_gal"]), afue_pct=int(round(p["afue"] * 100)), cord_mmbtu=p["cord_btu"] / 1e6,
                             pellet_mmbtu=p["pellet_btu"] / 1e6, tons_per_cord=p["tpc"], secondary=p["secondary"],
                             nox=p["ef_lp"]["Commercial/institutional propane combustion", "NOx"], bag_lb=40,
                             main_out=whole(G.main_households(p)))
    figs["steps"] = dict(appliances=sum(v["n"] for v in inv["cells"].values()), comp_lo=round(lo_c, 3),
                         comp_hi=round(hi_c, 3), drop_lo=whole(100 * lo_t), drop_hi=whole(100 * hi_t),
                         pre_gap=math.ceil(gap * 100) / 100.0)
    return figs
