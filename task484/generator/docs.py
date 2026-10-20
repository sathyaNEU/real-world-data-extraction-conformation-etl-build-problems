"""The pack's documents and workbooks: the tier rules (PDF), the data notes (txt), the census base, the ACS occupancy,
the county permit workbook and the building type code table (xlsx)."""
import numpy as np
from reportlab import rl_config
rl_config.invariant = 1
from reportlab.lib import colors                                    # noqa: E402
from reportlab.lib.pagesizes import letter                          # noqa: E402
from reportlab.lib.styles import ParagraphStyle                     # noqa: E402
from reportlab.lib.units import inch                                # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate        # noqa: E402
from openpyxl import Workbook                                       # noqa: E402
from openpyxl.styles import Font                                    # noqa: E402

import world as W   # noqa: E402

SECTIONS = {"cities": "2.1", "population": "2.2", "tier": "2.3", "permits": "3.1"}
BODY = ParagraphStyle("b", fontName="Times-Roman", fontSize=10.5, leading=13.5, spaceAfter=6)
HEAD = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=11, leading=14, spaceBefore=8, spaceAfter=4)
TITLE = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=15, leading=19, spaceAfter=4)
SUB = ParagraphStyle("s", fontName="Helvetica", fontSize=9, leading=12, textColor=colors.HexColor("#444444"),
                     spaceAfter=10)


def write_rules(path):
    title = "Municipal Revenue Sharing: 2027 Tier Determination"
    org = W.OFFICE
    d = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.95 * inch, rightMargin=0.95 * inch, topMargin=0.9 * inch,
                          bottomMargin=0.9 * inch, title=title, author=org, creator=org, producer=org, subject=title)

    def frame(cv, d_):
        cv.saveState()
        cv.setFont("Helvetica", 8)
        cv.setFillColor(colors.HexColor("#555555"))
        cv.drawString(0.95 * inch, 0.55 * inch, "%s  |  %s" % (org, title))
        cv.drawRightString(letter[0] - 0.95 * inch, 0.55 * inch, "Page %d" % d_.page)
        cv.restoreState()
    s = [Paragraph("Municipal Revenue Sharing, 2027 Distribution Year: Tier Determination", TITLE),
         Paragraph("%s. Adopted by the Joint Committee on Local Government Finance %s." % (W.OFFICE, W.RULES_DATE), SUB),
         Paragraph("1. Purpose", HEAD),
         Paragraph("Cities of 50,000 residents or more receive the upper-tier per-capita rate of the 2027 municipal "
                   "revenue-sharing distribution. The Population and Demographics Unit determines which cities near the "
                   "line stand at or above it on 1 April 2026.", BODY),
         Paragraph("2. Determination", HEAD),
         Paragraph("2.1 <b>Cities.</b> The twelve cities in the census housing base file.", BODY),
         Paragraph("2.2 <b>Population.</b> A city's population on 1 April 2026 is the persons living in its occupied "
                   "housing units plus its group quarters residents, estimated by the housing unit method: the city's "
                   "occupied housing units on 1 April 2026 times its average household size in the 2020 census, plus the "
                   "group quarters residents the state's 2026 group quarters survey counts. The city's housing units are "
                   "its 2020 census housing units, plus the units of buildings that received a certificate of occupancy "
                   "from 1 April 2020 to 31 March 2026, less the units demolished in that time.", BODY),
         Paragraph("2.3 <b>Tier.</b> A city whose population on 1 April 2026 is 50,000 or more is in the upper tier "
                   "for 2027.", BODY),
         Paragraph("3. Housing authorized", HEAD),
         Paragraph("3.1 For the Committee's review of housing growth beside the tiers: for each city, the housing units "
                   "authorized by building permits first issued in calendar 2025, single-family and multifamily. A "
                   "single-family unit is a building of one dwelling unit, detached, or attached to others by a wall "
                   "from ground to roof; every other new dwelling unit is multifamily. A building counts in the year its "
                   "first permit was issued.", BODY),
         Paragraph("4. Timetable", HEAD),
         Paragraph("The Unit reports the tier list to the Committee at its October meeting.", BODY)]
    d.build(s, onFirstPage=frame, onLaterPages=frame)


def write_notes(path, n):
    F = W.FILES
    lines = ["Population and Demographics Unit: data notes, 2026 city estimates",
             "Compiled 2026-10-14.",
             "",
             "%s" % F["census"],
             "  2020 census housing units, occupied units, household and group quarters population and average",
             "  household size, by city.",
             "",
             "%s" % F["acs"],
             "  American Community Survey 2020-2024 five-year estimates of total and occupied housing units, by city,",
             "  with margins of error (90 per cent).",
             "",
             "%s (%d rows), %s and %s (%d rows)" % (F["metro"], n["metro"], F["county"], F["civic"], n["civic"]),
             "  New residential building permits from the three permit systems the cities use, as each exports",
             "  them: one row per permit, with the issue date, status and, once issued, the certificate of",
             "  occupancy (final) date. Building type codes are listed in %s." % F["codes"],
             "",
             "%s (%d rows)" % (F["demo"], n["demo"]),
             "  Residential demolition permits finaled since the 2020 census.",
             "",
             "%s (%d rows)" % (F["gq"], n["gq"]),
             "  The state's group quarters survey: residents of each facility on 1 April 2026.",
             "",
             "%s (%d rows)" % (F["utility"], n["utility"]),
             "  %s residential accounts in the five cities it serves: one row per multi-unit premise" % W.UTILITY,
             "  per bill month (January 2021 to March 2026), with its residential meters and the accounts",
             "  active in the month.",
             "",
             "%s (%d rows), %s (%d rows)" % (F["schools"], n["schools"], F["voters"], n["voters"]),
             "  School enrollment on 1 October 2025 and voter registration on 1 September 2026, by city.",
             "",
             "Contact: Population and Demographics Unit."]
    open(path, "w", encoding="utf-8", newline="\r\n").write("\n".join(lines) + "\n")


def write_census(path, x):
    wb = Workbook()
    ws = wb.active
    ws.title = "2020 base"
    ws.append(["2020 Census: housing and population base, cities near the 50,000 line"])
    ws.append([])
    ws.append(["City", "Housing units", "Occupied units", "Household population", "Group quarters population",
               "Average household size"])
    for c in ws[3]:
        c.font = Font(bold=True)
    for c, v in sorted(x.items()):
        occd = int(round(v["hu"] * v["occ"]))
        hh = int(round(occd * v["pph"]))
        ws.append([c, v["hu"], occd, hh, int(round(v["gq"] * .97)), v["pph"]])
    for col, w in zip("ABCDEF", (16, 14, 15, 21, 25, 22)):
        ws.column_dimensions[col].width = w
    wb.save(path)


def write_acs(path, x, acs):
    rng = np.random.default_rng(W.SEED + 31)
    wb = Workbook()
    ws = wb.active
    ws.title = "ACS 2020-2024"
    ws.append(["American Community Survey 2020-2024 five-year estimates: housing occupancy by city"])
    ws.append([])
    ws.append(["City", "Total housing units", "MOE", "Occupied housing units", "MOE", "Vacant housing units", "MOE"])
    for c in ws[3]:
        c.font = Font(bold=True)
    for c, v in sorted(x.items()):
        tot = int(round(v["hu"] * 1.07))
        occd = int(round(tot * acs[c]))
        ws.append([c, tot, int(rng.integers(300, 700)), occd, int(rng.integers(500, 900)), tot - occd,
                   int(rng.integers(250, 500))])
    wb.save(path)


def write_county(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Permits"
    ws.append(["County Building Department: new residential permits, municipalities served"])
    ws.append([])
    ws.append(["Permit #", "Municipality", "APN", "Issued", "Status", "Type Code", "Units", "Project", "C of O Date",
               "Expired"])
    for c in ws[3]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([v if v != "" else None for v in r])
        for cell in ws[ws.max_row]:
            if hasattr(cell.value, "year"):
                cell.number_format = "mm/dd/yyyy"
    wb.save(path)


def write_codes(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Building types"
    ws.append(["System", "Code", "Description", "Structure category", "In use from", "In use to"])
    for c in ws[1]:
        c.font = Font(bold=True)
    rows = [("Metro permits", "SFD", "Single-family detached", "Single-family", "2015-07-01", ""),
            ("Metro permits", "SFA", "Single-family attached (townhouse)", "Single-family", "2015-07-01", ""),
            ("Metro permits", "MF2-4", "Two to four units", "Multifamily", "2015-07-01", ""),
            ("Metro permits", "MF5+", "Five or more units", "Multifamily", "2015-07-01", ""),
            ("County building department", "101", "One-family detached", "Single-family", "2009-01-01", ""),
            ("County building department", "102", "One-family attached", "Single-family", "2009-01-01", ""),
            ("County building department", "103", "Two-family", "Multifamily", "2009-01-01", ""),
            ("County building department", "104", "Three and four family", "Multifamily", "2009-01-01", ""),
            ("County building department", "105", "Five or more family", "Multifamily", "2009-01-01", ""),
            ("Civic permits", "R3", "Residential, one or two dwelling units", "See note", "2012-04-01", "2024-12-31"),
            ("Civic permits", "R3T", "Residential, townhouse", "Single-family", "2012-04-01", "2024-12-31"),
            ("Civic permits", "R2", "Residential, three or more dwelling units", "Multifamily", "2012-04-01", "2024-12-31"),
            ("Civic permits", "1-DET", "One unit, detached", "Single-family", "2025-01-01", ""),
            ("Civic permits", "1-ATT", "One unit, attached", "Single-family", "2025-01-01", ""),
            ("Civic permits", "2-UNIT", "Two units", "Multifamily", "2025-01-01", ""),
            ("Civic permits", "MULTI", "Three or more units", "Multifamily", "2025-01-01", "")]
    for r in rows:
        ws.append([v if v != "" else None for v in r])
    ws.append([])
    ws.append(["Note: R3 covered one- and two-unit buildings; the units field separates them."])
    for col, w in zip("ABCDEF", (28, 8, 40, 20, 12, 12)):
        ws.column_dimensions[col].width = w
    wb.save(path)
