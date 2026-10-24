"""Writes the shipped files from the world's objects, each in the layout and the vocabulary of the system that keeps it."""
import csv
import datetime as dt

import numpy as np
import openpyxl
import reportlab.rl_config
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import world as W

reportlab.rl_config.invariant = 1
REG_CODE = dict(UNC="WS-NC", CNC="WS-EPA", CAT="WS-CAT", PEL="PS", FPL="FP", HYD="OWB")
SURVEY_LABEL = dict(UNC="Wood stove, older or not certified", CNC="Wood stove, EPA certified (no catalyst)",
                    CAT="Wood stove, EPA certified with catalyst", PEL="Pellet stove", FPL="Open fireplace",
                    HYD="Outdoor wood boiler")
APP_LABEL = dict(UNC="Non-certified wood stove", HYD="Outdoor wood boiler", CNC="EPA certified wood stove",
                 PEL="Pellet stove", HP="Ductless heat pump")
AREA_NAME = {a[0]: a[1] for a in W.AREAS}
TOOL_SHARE = {True: dict(UNC=0.30, CNC=0.32, CAT=0.08, PEL=0.20, FPL=0.0, HYD=0.10),
              False: dict(UNC=0.18, CNC=0.22, CAT=0.05, PEL=0.08, FPL=0.47, HYD=0.0)}
TOOL_USE = {True: dict(UNC=3.9, CNC=3.6, CAT=3.6, PEL=3.0, FPL=0.9, HYD=7.0),
            False: dict(UNC=1.0, CNC=0.9, CAT=0.9, PEL=0.8, FPL=0.3, HYD=0.0)}
TOOL_SECONDARY = 2.4


def iso(d):
    return d.isoformat() if d else ""


def write_csv(path, header, rows):
    with open(path, "w", newline="") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(header)
        w.writerows(rows)


def area_2020(d):
    return "" if d.moved else d.area


def registry_rows(devs):
    """One row per registration. A parcel split closes the appliance's record and registers it again under a new id
    and the new parcel on the same day, the install year carried over; nothing links the two records. Some closures
    for other reasons are also recorded only as a closed record."""
    r = np.random.default_rng(W.SEED + 51)
    rows = []
    nxt = [900000]
    for d in devs:
        closed = d.split if d.split else d.removed
        reason = "Record closed" if d.split else (d.reason or "")
        if d.reason == "Owner removed" and r.random() < 0.5:
            reason = "Record closed"
        rows.append([d.did, d.parcel, d.address, AREA_NAME[d.area], d.county, REG_CODE[d.cl], d.installed,
                     iso(d.registered), "C" if closed else "A", iso(closed), reason, d.replaces or ""])
        if d.split:
            nxt[0] += int(r.integers(1, 9))
            rows.append(["WD-%06d" % nxt[0], W.SPLIT_SUFFIX[d.county](d.parcel), d.address, AREA_NAME[d.area], d.county,
                         REG_CODE[d.cl], d.installed, iso(d.split), "C" if d.removed else "A", iso(d.removed),
                         d.reason or "", ""])
    rows.sort(key=lambda x: x[0])
    return rows


def registry_csv(path, devs):
    write_csv(path, ["device_no", "parcel_number", "site_address", "maintenance_area", "county", "appliance_code",
                     "year_installed", "registered_on", "status", "status_date", "status_reason", "replaces_device_no"],
              registry_rows(devs))


def survey_csv(path, resp):
    r = np.random.default_rng(W.SEED + 21)
    rows = []
    for i, x in enumerate(sorted(resp, key=lambda x: (x["dev"].area, x["dev"].did))):
        d = x["dev"]
        yr = str(d.installed) if r.random() < 0.7 else "'%02d" % (d.installed % 100)
        rows.append(["R%05d" % (10001 + i), AREA_NAME[d.area], SURVEY_LABEL[d.cl], yr, "Yes" if d.primary else "No",
                     ("%g" % x["amount"]), x["unit"], "%.4f" % x["weight"]])
    write_csv(path, ["respondent_id", "maintenance_area", "appliance", "year_installed", "wood_main_heat",
                     "amount_burned_2024_25", "amount_unit", "weight"], rows)


def fiscal(d):
    y = d.year if d.month >= 7 else d.year - 1
    return "FY%d-%02d" % (y, (y + 1) % 100)


def applications_csv(path, apps, devs):
    r = np.random.default_rng(W.SEED + 31)
    rows = []
    allrec = [d for d in devs if d.replaces is None and d.installed <= 2014]
    for a in apps:
        d = a["old"]
        rec = d.removed - dt.timedelta(days=int(r.integers(40, 160)))
        key_p = d.parcel if r.random() < 0.6 else ""
        heat = ("P" if d.primary else "S") if r.random() < 0.6 else ("main heat source" if d.primary else "supplemental")
        voucher = {"CNC": 2500, "PEL": 3000, "HP": 4500}[a["new"]] + 500 * int(r.integers(0, 3))
        rows.append([rec, key_p, d.address, d.county, APP_LABEL[d.cl], heat, APP_LABEL[a["new"]], iso(d.removed),
                     "Completed", voucher])
    for k in range(int(len(apps) * 0.12)):           # withdrawn and denied applications: no changeout
        d = allrec[int(r.integers(len(allrec)))]
        if d.programme or d.cl not in W.CHANGED:
            continue
        rec = dt.date(int(r.integers(2015, 2026)), int(r.integers(1, 13)), int(r.integers(1, 28)))
        rows.append([rec, d.parcel if r.random() < 0.6 else "", d.address, d.county, APP_LABEL[d.cl],
                     ("P" if d.primary else "S"), "", "", str(r.choice(["Withdrawn", "Denied"])), ""])
    rows.sort(key=lambda x: (x[0], x[2]))
    out = []
    for i, x in enumerate(rows):
        out.append(["CO-%s-%04d" % (str(x[0].year)[2:], i + 1), iso(x[0]), fiscal(x[0])] + x[1:])
    write_csv(path, ["application_no", "received", "fiscal_year", "parcel_number", "site_address", "county",
                     "existing_appliance", "heat", "replacement", "installed_on", "status", "voucher_usd"], out)


def crosswalk_csv(path, devs):
    seen = {}
    for d in devs:
        seen.setdefault(d.parcel, (AREA_NAME[d.area], "" if d.moved else AREA_NAME[d.area], d.county))
        if d.split:
            p = W.SPLIT_SUFFIX[d.county](d.parcel)
            seen.setdefault(p, (AREA_NAME[d.area], "" if d.moved else AREA_NAME[d.area], d.county))
    rows = [[p, c, a14, a20] for p, (a14, a20, c) in sorted(seen.items())]
    write_csv(path, ["parcel_number", "county", "area_2014_boundary", "area_2020_boundary"], rows)


def plan_workbook(path, base):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RWC base year 2014"
    ws.append(["Maintenance area", "County", "SCC", "Appliance", "Appliances", "Cords per appliance",
               "Wood burned (tons)", "PM2.5 factor (lb/ton)", "PM2.5 (tons)"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for code, county, cl, n in base:
        scc, desc, n0, prim, cp, cs, ef, ef14, dflt = W.CLASSES[cl]
        wood = W.tons_of(cl, n * dflt)
        ws.append([AREA_NAME[code], county, scc, desc, n, dflt, round(wood, 1), ef14, round(wood * ef14 / 2000, 2)])
    for i, w in enumerate([16, 10, 12, 50, 11, 18, 18, 20, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    nt = wb.create_sheet("Notes")
    for line in ["2014 base year, residential wood combustion, as adopted in each area's maintenance plan.",
                 "Appliances: the device registry extract of 31 March 2015, appliances installed by 2014 inside the "
                 "2014 maintenance area boundaries.",
                 "Activity: the 2014 regional defaults (cords per appliance; pellet stoves in tons of pellets), "
                 "%.3f tons per cord." % W.TONS_PER_CORD,
                 "Factors: the 2014 inventory tool's PM2.5 factors."]:
        nt.append([line])
    nt.column_dimensions["A"].width = 110
    wb.save(path)


def factors_csv(path):
    rows = []
    for cl in W.CL:
        scc, desc, *_x = W.CLASSES[cl]
        ef = W.CLASSES[cl][6]
        for pol, k in (("PM2.5", 1.0), ("PM10", 1.04), ("CO", 7.6), ("VOC", 1.9)):
            rows.append([scc, desc, pol, "%.1f" % (ef * k), "lb/ton", "2023"])
    write_csv(path, ["scc", "description", "pollutant", "factor", "unit", "edition"], rows)


def scc_crosswalk_csv(path):
    rows = []
    for cl in W.CL:
        rows.append([W.CLASSES[cl][0], W.CLASSES[cl][1], REG_CODE[cl], SURVEY_LABEL[cl], APP_LABEL.get(cl, "")])
    write_csv(path, ["scc", "description", "registry_appliance_code", "survey_appliance", "application_appliance"], rows)


def acs_csv(path, bgs, which):
    rows = []
    for b in bgs:
        total = b["hu"] - int(b["hu"] * 0.07)
        est = b["acs1923"] if which == "1923" else b["acs2024"]
        moe = b["moe1923"] if which == "1923" else b["moe2024"]
        rows.append(["1500000US" + b["geoid"], "Block Group %s, Census Tract %s, %s County, Washington" %
                     (b["geoid"][-1], str(int(b["geoid"][5:11]) / 100.0).rstrip("0").rstrip("."), b["county"]),
                     total, int(total * 0.05) + 12, est, moe])
    write_csv(path, ["GEO_ID", "NAME", "B25040_001E", "B25040_001M", "B25040_008E", "B25040_008M"], rows)


def assessor_csv(path, bgs):
    rows = []
    for b in bgs:
        hu_in = int(round(b["hu"] * b["inside"]))
        if hu_in:
            rows.append([b["geoid"], b["county"], "inside a maintenance area", hu_in])
        if b["hu"] - hu_in:
            rows.append([b["geoid"], b["county"], "", b["hu"] - hu_in])
    write_csv(path, ["block_group", "county", "maintenance_area_2020", "housing_units"], rows)


def distractor_csvs(t, devs):
    r = np.random.default_rng(W.SEED + 41)
    days = [W.SEASON[0] + dt.timedelta(days=i) for i in range((W.SEASON[1] - W.SEASON[0]).days + 1)]
    write_csv(t("hdd"), ["station", "date", "heating_degree_days_f"],
              [[s, iso(d), max(0, int(round(65 - (38 + 18 * np.cos((d - dt.date(2026, 1, 15)).days / 58.0)
                                                      + r.normal(0, 6)))))] for s in ("KSLW", "KCRN", "KTMS") for d in days])
    rows = []
    for code in W.CODES:
        for d in days:
            if r.random() < 0.07:
                rows.append([AREA_NAME[code], iso(d), str(r.choice(["Stage 1", "Stage 2"])), "07:00", "next day 07:00"])
    write_csv(t("curtail"), ["maintenance_area", "date", "stage", "from", "until"], rows)
    write_csv(t("monitor"), ["site", "date", "pm25_24h_ug_m3"],
              [[AREA_NAME[c] + " FRM", iso(d), "%.1f" % max(1.0, r.gamma(2.0, 6.0))] for c in W.CODES for d in days])
    write_csv(t("permits"), ["ranger_district", "season", "permits_sold", "cords_permitted"],
              [[dname, s, int(r.integers(300, 1500)), int(r.integers(1500, 6000))] for dname in ("Selwick", "Corrin", "Tamsey")
               for s in ("2022-23", "2023-24", "2024-25", "2025-26")])


STY = getSampleStyleSheet()
BODY = ParagraphStyle("b", parent=STY["BodyText"], fontName="Helvetica", fontSize=9.5, leading=12.5, spaceAfter=6)
H = ParagraphStyle("h", parent=STY["Heading2"], fontName="Helvetica-Bold", fontSize=11, spaceBefore=8, spaceAfter=4)
T = ParagraphStyle("t", parent=STY["Title"], fontName="Helvetica-Bold", fontSize=14, spaceAfter=10)


def _pdf(path, title, story, author):
    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.8 * inch, rightMargin=0.8 * inch, topMargin=0.7 * inch,
                            bottomMargin=0.7 * inch, title=title, author=author, creator=author, subject=title,
                            producer=W.AGENCY)

    def footer(c, d):
        c.saveState()
        c.setFont("Helvetica", 7.5)
        c.drawString(0.8 * inch, 0.45 * inch, W.AGENCY + "  |  " + title)
        c.drawRightString(letter[0] - 0.8 * inch, 0.45 * inch, "Page %d" % d.page)
        c.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def _table(rows, widths=None):
    t = Table(rows, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 8.5), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.5),
                           ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.black), ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 2), ("TOPPADDING", (0, 0), (-1, -1), 2)]))
    return t


def provision_pdf(path):
    s = [Paragraph("PM2.5 Maintenance Plans: Contingency Provisions for Residential Wood Combustion", T),
         Paragraph("Adopted by the Board of Directors on 9 December 2016 and incorporated in the maintenance plans for the "
                   "six maintenance areas. Section 4 is reproduced as adopted.", BODY),
         Paragraph("4.1 Trigger", H),
         Paragraph("The contingency measure for a maintenance area takes effect for the next heating season when the "
                   "periodic inventory for the latest complete heating season shows the area's PM2.5 emissions from "
                   "residential wood combustion reduced from the 2014 base year by less than the area's commitment in "
                   "Table 4-1.", BODY),
         Paragraph("4.2 Base year", H),
         Paragraph("The 2014 base year is recalculated on the methods of the periodic inventory it is compared with: the "
                   "wood each appliance in the plan's 2014 base year count burned, on the Agency's household survey measure "
                   "of wood burned, times the emission factors the periodic inventory uses. The appliance count is not "
                   "revised.", BODY),
         Paragraph("4.3 Areas", H),
         Paragraph("A maintenance area is read on its boundary as revised by the Department of Ecology in 2020.", BODY),
         Paragraph("4.4 Measure", H),
         Paragraph("The contingency measure lowers the Stage 1 curtailment forecast threshold from 30 to 25 micrograms per "
                   "cubic meter and prohibits the use of non-certified wood stoves during any declared Stage 1 episode.", BODY),
         Spacer(1, 6), Paragraph("Table 4-1. Commitments", H),
         _table([["Maintenance area", "County", "Reduction from the 2014 base year (per cent)"]] +
                [[a[1], a[2], "%d" % a[6]] for a in W.AREAS], [2.0 * inch, 1.4 * inch, 3.0 * inch])]
    _pdf(path, "Contingency provisions", s, W.AGENCY)


def methods_pdf(path):
    s = [Paragraph("Periodic Inventory of Residential Wood Combustion: Methods Appendix, 2025-26 Season", T),
         Paragraph("Prepared by %s, Air Quality Planning. Applies to the 2025-26 periodic inventory." % W.PEOPLE["inventory"], BODY),
         Paragraph("A.1 Appliances", H),
         Paragraph("The appliances in an inventory season are the registry's wood-burning devices in place in the season "
                   "(1 October to 30 April), each device once, on the 2020 maintenance area boundaries (the parcel "
                   "crosswalk). Device types are mapped to SCCs by the SCC crosswalk.", BODY),
         Paragraph("A.2 Activity", H),
         Paragraph("Wood burned per appliance comes from the Agency's household survey of registered appliances, the "
                   "2024-25 survey for this inventory, estimated for each maintenance area from that area's respondents, "
                   "by appliance class, with each respondent's weight. The survey also records whether wood is the "
                   "household's main heat source. Reported amounts are converted to cords (a face cord is one third of a "
                   "cord, a pickup load half a cord) and to tons at %.3f tons per cord; pellets are reported in tons or in "
                   "%d-pound bags. An appliance in place but not used in the season counts at zero. Activity is not "
                   "adjusted for weather." % (W.TONS_PER_CORD, int(W.BAG_LB)), BODY),
         Paragraph("A.3 Emission factors", H),
         Paragraph("Pounds of pollutant per ton of wood burned, from the 2023 edition of the factor table, by SCC.", BODY),
         Paragraph("A.4 Outside the maintenance areas", H),
         Paragraph("The counties' households outside the maintenance areas are inventoried on the state's residential "
                   "wood combustion tool (see the tool's method note), each block group by its housing units outside the "
                   "areas.", BODY)]
    _pdf(path, "Methods appendix", s, W.PEOPLE["inventory"])


def tool_pdf(path):
    rows = [["Appliance", "Share of main-heat households", "Share of secondary households", "Use, main heat",
             "Use, secondary"]]
    for cl in W.CL:
        rows.append([W.CLASSES[cl][1].replace("Woodstove: ", ""), "%.2f" % TOOL_SHARE[True][cl], "%.2f" % TOOL_SHARE[False][cl],
                     "%.1f" % TOOL_USE[True][cl], "%.1f" % TOOL_USE[False][cl]])
    s = [Paragraph("State Residential Wood Combustion Tool: Method Note", T),
         Paragraph("Households heating mainly with wood are taken from the American Community Survey table B25040 (house "
                   "heating fuel: wood) for each block group, from the five-year release whose final year is two years "
                   "before the first year of the inventory season. Households burning wood as a secondary heat source are "
                   "%.1f times the main-heat households. Each household group is split across appliances by the shares "
                   "below and burns the use shown (cords a year; pellet stoves in tons of pellets)." % TOOL_SECONDARY, BODY),
         _table(rows, [2.4 * inch, 1.1 * inch, 1.1 * inch, 0.9 * inch, 0.9 * inch])]
    _pdf(path, "Tool method note", s, "Department of Ecology")
