"""Reads the cut pack back and runs the estimators: the cities' permits from the three systems as one table, the
housing units on the estimate date, the lease-up curve from the utility's accounts, and the population under each
route; the asks' permit counts."""
import csv
import datetime as dt
import os

import numpy as np
from openpyxl import load_workbook

import world as W

EST = W.EST


def find(folder, stem):
    for fn in sorted(os.listdir(folder)):
        if fn.startswith(stem):
            return os.path.join(folder, fn)
    raise FileNotFoundError(stem)


def _rows(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def _us(s):
    if not s:
        return None
    m, d, y = s.split("/")
    return dt.date(int(y), int(m), int(d))


def _iso(s):
    return dt.date.fromisoformat(s) if s else None


def codes(target):
    ws = load_workbook(find(target, "building_type_code_table"), read_only=True).active
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r and r[1] is not None and r[0]:
            out[(r[0].split()[0].lower(), str(r[1]))] = r[3]
    return out


def permits(target):
    """[dict(city, id, parcel, project, issued, status, category, units, co)]: one row per permit, category
    'Single-family' or 'Multifamily' read from the code table (an R3 by its units)."""
    cat = codes(target)
    out = []
    for r in _rows(find(target, "metro_permits")):
        out.append(dict(city=r["Jurisdiction"].title(), id=r["PermitNumber"], parcel=r["ParcelID"], project=r["ProjectName"],
                        issued=_us(r["IssueDate"]), status=r["Status"].lower(), category=cat[("metro", r["BldgType"])],
                        units=int(r["DwellingUnits"]), co=_us(r["FinalDate"])))
    for r in _rows(find(target, "civic_permit_records")):
        c = cat[("civic", r["occupancy_code"])]
        if c == "See note":
            c = "Single-family" if int(r["units"]) == 1 else "Multifamily"
        out.append(dict(city=r["city"], id=r["record_id"], parcel=r["parcel_number"], project=r["project_name"],
                        issued=_iso(r["issued_on"]), status=r["status"].lower(), category=c, units=int(r["units"]),
                        co=_iso(r["co_issued_on"])))
    ws = load_workbook(find(target, "county_building_permits"), read_only=True).active
    for r in ws.iter_rows(min_row=4, values_only=True):
        if not r or not r[0]:
            continue
        d = lambda v: v.date() if hasattr(v, "date") else v
        out.append(dict(city=r[1], id=r[0], parcel=r[2], project=r[7] or "", issued=d(r[3]), status=str(r[4]).lower(),
                        category=cat[("county", str(r[5]))], units=int(r[6]), co=d(r[8])))
    return out


def census(target):
    ws = load_workbook(find(target, "census_2020_city_housing_base"), read_only=True).active
    return {r[0]: dict(hu=r[1], occupied=r[2], hhpop=r[3], gq=r[4], pph=r[5])
            for r in ws.iter_rows(min_row=4, values_only=True) if r and r[0]}


def acs(target):
    ws = load_workbook(find(target, "acs_2020_2024_city_occupancy"), read_only=True).active
    return {r[0]: r[3] / r[1] for r in ws.iter_rows(min_row=4, values_only=True) if r and r[0]}


def demolitions(target):
    out = {}
    for r in _rows(find(target, "demolition_permits")):
        d = _iso(r["finaled"])
        if W.CENSUS <= d < EST:
            out[r["city"]] = out.get(r["city"], 0) + int(r["units_removed"])
    return out


def gq(target):
    out = {}
    for r in _rows(find(target, "group_quarters_survey")):
        out[r["city"]] = out.get(r["city"], 0) + int(r["residents_apr1_2026"])
    return out


def utility(target):
    """{premise: (city, {bill month: (meters, accounts)})}."""
    out = {}
    for r in _rows(find(target, "nvp_residential_accounts")):
        y, m = r["bill_month"].split("-")
        p = out.setdefault(r["premise_id"], (r["service_city"].title(), {}))
        p[1][dt.date(int(y), int(m), 1)] = (int(r["residential_meters"]), int(r["active_residential_accounts"]))
    return out


def _mindex(a, b):
    return (b.year - a.year) * 12 + b.month - a.month


def curve(target, form="pooled", until=None):
    """{months since the first bill: share of meters with an active account} over the premises whose first bill
    came after January 2021 (buildings opened in the record). pooled: accounts over meters at each month; mean: the
    mean of the buildings' shares; exp: a fitted level times one less exp(-m/tau). until: only bills to that month."""
    num, den, shares = {}, {}, {}
    for city, bills in utility(target).values():
        first = min(bills)
        if first <= dt.date(2021, 1, 1):
            continue
        for bm, (m, a) in bills.items():
            if until and bm > until:
                continue
            k = _mindex(first, bm)
            num[k] = num.get(k, 0) + a
            den[k] = den.get(k, 0) + m
            shares.setdefault(k, []).append(a / m)
    ks = sorted(k for k in den if len(shares[k]) >= 5)
    if form == "pooled":
        c = {k: num[k] / den[k] for k in ks}
    elif form == "mean":
        c = {k: float(np.mean(shares[k])) for k in ks}
    else:
        x = np.array(ks, float)
        y = np.array([num[k] / den[k] for k in ks])
        best = None
        for lev in np.arange(.90, .99, .0025):
            for tau in np.arange(3.0, 16.0, .05):
                e = float(((y - lev * (1 - np.exp(-(x + .5) / tau))) ** 2).sum())
                if best is None or e < best[0]:
                    best = (e, lev, tau)
        c = {k: best[1] * (1 - np.exp(-(k + .5) / best[2])) for k in range(0, 80)}
    return c


def housing(target):
    """{city: (stock units at the stock's occupancy, [(multifamily units, certificate-of-occupancy date)])} on the
    estimate date: the 2020 units less demolitions plus every single-family completion, and the multifamily
    buildings completed since the census."""
    base, dem = census(target), demolitions(target)
    stock = {c: base[c]["hu"] - dem.get(c, 0) for c in base}
    mf = {c: [] for c in base}
    for p in permits(target):
        if p["co"] and W.CENSUS <= p["co"] < EST:
            if p["category"] == "Multifamily":
                mf[p["city"]].append((p["units"], p["co"]))
            else:
                stock[p["city"]] += p["units"]
    return {c: (stock[c], mf[c]) for c in base}


def population(target, route="golden", form="pooled", direct=False, since=None):
    """{city: residents on 1 April 2026}. golden: the stock at the city's ACS occupancy, each multifamily building
    completed since the census at the lease-up curve's share for its months from opening to April 2026;
    textbook: every unit at the ACS occupancy; census: every unit at the 2020 census occupancy. direct: the covered
    cities' buildings at their March 2026 accounts. since: the curve only for buildings opened on or after it."""
    base, occ, gqs, h = census(target), acs(target), gq(target), housing(target)
    cv = curve(target, form) if route == "golden" else None
    if cv:
        tail = float(np.mean([v for k, v in cv.items() if k >= max(cv) - 6]))
    util = utility(target) if direct else {}
    opened = {}
    for pid, (city, bills) in util.items():
        first = min(bills)
        if first > dt.date(2021, 1, 1):
            m, a = bills[max(bills)]
            opened.setdefault(city, {}).setdefault((m, first), []).append(a)
    out = {}
    for c in base:
        stock, mf = h[c]
        rate = base[c]["occupied"] / base[c]["hu"] if route == "census" else occ[c]
        if route in ("textbook", "census"):
            occd = (stock + sum(u for u, _ in mf)) * rate
        else:
            occd = stock * rate
            for u, co in mf:
                k = _mindex(dt.date(co.year, co.month, 1), dt.date(2026, 4, 1))
                if since and co < since:
                    occd += u * rate
                    continue
                key = (u, dt.date(co.year, co.month, 1))
                if direct and c in opened and key in opened[c] and opened[c][key]:
                    occd += opened[c][key].pop(0)
                    continue
                occd += u * (cv[k] if k in cv else tail)
        out[c] = occd * base[c]["pph"] + gqs[c]
    return out


def tier(pop):
    return sorted(c for c, v in pop.items() if v >= W.LINE)


# the asks: housing authorized in 2025
def authorized(target, route="golden"):
    """{city: (single-family units, multifamily units)} on permits first issued in 2025. golden: a building's units
    in the year of its first permit (a permit for a parcel and project an earlier permit already covered is not
    counted again); category from the code table. rows: every 2025 permit; units: category by units (one unit
    single-family)."""
    ps = permits(target)
    first = {}
    for p in ps:
        key = (p["city"], p["parcel"], p["project"], p["units"])
        if key not in first or p["issued"] < first[key]:
            first[key] = p["issued"]
    out = {c: [0, 0] for c in W.CITIES}
    for p in ps:
        if p["issued"].year != 2025:
            continue
        if route == "golden" and first[(p["city"], p["parcel"], p["project"], p["units"])] < p["issued"]:
            continue
        sf = (p["units"] == 1) if route == "units" else (p["category"] == "Single-family")
        out[p["city"]][0 if sf else 1] += p["units"]
    return {c: tuple(v) for c, v in out.items()}
