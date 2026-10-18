"""task246 goldens: the measurement review workbook (each zone's figure from the readings, from its children's own arm
circumference with the range over the standard variants, the note, and the household figures for the annex; each
anthropometrist's standardisation results) and the one-page note to the bulletin desk with its chart."""
import io
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer

import params as P
import postprocess

XLSX = "lr2027_muac_measurement_review.xlsx"
PDF = "lr2027_gam_measurement_note.pdf"
ORG = "%s, %s" % (P.COUNTY_GOV, P.UNIT)
AUTHOR = "Winny Karisa"
BAND_COLOURS = ["#E8F1E4", "#FBF1D3", "#F8DEC6", "#F1C9C4"]


def _head(ws, row, head, height=58):
    thin = Side(style="thin", color="A6ADB8")
    for j, h in enumerate(head, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(bold=True, size=9.5)
        c.fill = PatternFill("solid", fgColor="DDE4EE")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = Border(bottom=thin)
    ws.row_dimensions[row].height = height


def _workbook(path, A, test):
    wb = Workbook()
    ws = wb.active
    ws.title = "Zones"
    ws["A1"] = "MUAC measurement review, 2027 long-rains round, with the household figures for the bulletin's annex"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Laga County, twelve survey zones, children 6 to 59 months. Figures and the measurement note as defined in "
                "the surveillance protocol v2.1, sections 6, 7, 9 and 10.")
    ws["A2"].font = Font(italic=True, size=9)
    head = ["Zone", "Sub-county", "Children measured", "Of which oedema", "Acute malnutrition, MUAC readings (%)",
            "Acute malnutrition, children's own arm circumference (%)", "Own arm circumference, accepted range (%)",
            "Measurement note", "Figure for the bulletin (%)", "Alert band", "Household response rate (%)",
            "Crude death rate (deaths/10,000/day)", "Under-five death rate (deaths/10,000/day)"]
    _head(ws, 4, head, 66)
    for i, r in enumerate(A["zones"], 5):
        vals = [r["name"], r["sub_county"], r["children"], r["oedema"], r["readings"], r["own"],
                "%.1f to %.1f" % tuple(r["own_range"]), "Note" if r["note"] else "", r["published"], r["band"],
                r["response"], r["cdr"], r["u5dr"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(size=10, bold=(j == 8 and r["note"]))
            if j in (5, 6, 9, 11):
                c.number_format = "0.0"
            elif j in (12, 13):
                c.number_format = "0.00"
            elif j in (3, 4):
                c.number_format = "0"
            c.alignment = Alignment(horizontal="left" if j in (1, 2, 10) else ("center" if j in (7, 8) else "right"))
            if r["note"]:
                c.fill = PatternFill("solid", fgColor="FBEFD9")
    n = 5 + len(A["zones"])
    notes = [
        "Acute malnutrition: children aged 6 to 59 months measured in the zone's clusters, bilateral oedema counted, the "
        "rest's share below 125 mm read from the normal distribution with their mean and standard deviation (section 6). "
        "Children's own arm circumference: each reading less its anthropometrist's offset against the supervisor, the "
        "variance less the mean of the anthropometrists' variance of departures from the supervisor about their offset "
        "(sheet Standardisation). Accepted range: the same figure under the standard variants (the correction by zone, "
        "the sessions apart, divisors n, the supervisor's own repeat error, only offsets that stand out, the error "
        "through the test children's means); every variant gives the same notes.",
        "Measurement note: the two figures, each to one decimal, differ by 2.0 points or more (section 7); a noted zone "
        "goes into the bulletin at its children's own figure. Alert bands: Normal under 5.0, Alert 5.0 to 9.9, Alarm "
        "10.0 to 14.9, Emergency 15.0 or more (section 8).",
        "Household figures (sections 9 and 10), from the county's household form: each household's last visit; before "
        "21 June the form coded an absent household 3 and an empty dwelling 4, and a member who died away from home 3 "
        "and one who moved with the mother or caregiver 4 (the dictionary gives version 4's codes, the other way "
        "round). Deaths include members who died away from home and newborns on the births module; each member once.",
        "Sources: NCHA and PRK anthropometry exports and standardisation sheets; survey teams workbook; cluster sample; "
        "county household form exports (visits, member rosters, births)."]
    for k, t in enumerate(notes):
        c = ws.cell(row=n + 1 + k, column=1, value=t)
        c.font = Font(size=8.5, italic=(k == len(notes) - 1))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=n + 1 + k, start_column=1, end_row=n + 1 + k, end_column=len(head))
        ws.row_dimensions[n + 1 + k].height = [62, 30, 44, 18][k]
    for col, w in zip("ABCDEFGHIJKLM", [20, 13, 10, 9, 13, 14, 13, 11, 11, 10, 11, 12, 12]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B5"
    ws.print_title_rows = "4:4"

    ts = wb.create_sheet("Standardisation")
    ts["A1"] = "Each anthropometrist against the supervisor, standardisation sessions of 10 and 28 June 2027"
    ts["A1"].font = Font(bold=True, size=11)
    _head(ts, 3, ["Anthropometrist", "Readings", "Offset against the supervisor (mm)",
                  "Departure from the supervisor about the offset, SD (mm)", "Repeat error, TEM (mm)"], 58)
    for i, r in enumerate(test, 4):
        vals = [r["name"], r["readings"], r["offset"], r["departure_sd"], r["tem"]]
        for j, v in enumerate(vals, 1):
            c = ts.cell(row=i, column=j, value=v)
            c.number_format = "0" if j == 2 else "0.0"
            c.alignment = Alignment(horizontal="left" if j == 1 else "right")
    m = 4 + len(test)
    c = ts.cell(row=m + 1, column=1, value="Reference: the supervisor's mean of her two readings of each test child "
                                           "(section 5); an anthropometrist's sessions taken together.")
    c.font = Font(size=8.5, italic=True)
    ts.merge_cells(start_row=m + 1, start_column=1, end_row=m + 1, end_column=5)
    for col, w in zip("ABCDE", [20, 10, 16, 20, 14]):
        ts.column_dimensions[col].width = w
    wb.save(path)
    postprocess.scrub(path, AUTHOR, "2027-07-22")
    postprocess.finish_office(path, "2027-07-19T06:10:00Z", "2027-07-22T13:40:00Z", ORG)


def chart_png(A):
    rows = A["zones"]
    fig, ax = plt.subplots(figsize=(6.8, 4.5), dpi=200)
    edges = [lo for lo, _n in P.BANDS] + [24.0]
    for (lo, name), hi, col in zip(P.BANDS, edges[1:], BAND_COLOURS):
        ax.axvspan(lo, hi, color=col, zorder=0, lw=0)
        ax.text((lo + hi) / 2, len(rows) - 0.25, name, ha="center", va="center", fontsize=7.2, color="#4A4A4A")
    y = list(range(len(rows)))[::-1]
    for yy, r in zip(y, rows):
        col = "#9A3B21" if r["note"] else "#5F6F86"
        ax.plot([r["readings"], r["own"]], [yy, yy], color=col, lw=1.5, zorder=2)
        ax.scatter([r["readings"]], [yy], s=22, facecolor="white", edgecolor=col, lw=1.2, zorder=3)
        ax.scatter([r["own"]], [yy], s=24, color=col, zorder=4)
        if r["note"]:
            ax.text(max(r["readings"], r["own"]) + 0.35, yy, "Note", va="center", fontsize=7, color="#9A3B21",
                    fontweight="bold")
    ax.set_yticks(y)
    ax.set_yticklabels([r["name"] for r in rows], fontsize=7.6)
    ax.set_xlim(0, 21)
    ax.set_ylim(-0.7, len(rows) + 0.15)
    ax.set_xlabel("Acute malnutrition by MUAC, children 6 to 59 months, per cent\n(open: from the readings; filled: "
                  "from the children's own arm circumference)", fontsize=7.0)
    ax.set_title(A["chart_title"], fontsize=9.3, loc="left", fontweight="bold")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.tick_params(axis="x", labelsize=7.5)
    fig.text(0.01, 0.012, "Note: the two figures differ by 2.0 points or more (protocol v2.1, section 7). Bands: section 8."
                          "\nSource: NCHA and PRK anthropometry and standardisation, long-rains round 2027.",
             fontsize=5.8, color="#555555", va="bottom")
    fig.tight_layout(rect=(0, 0.06, 1, 1))
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def _note(path, A):
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=9.2, leading=12.6, spaceAfter=5)
    head = ParagraphStyle("h", parent=body, fontName="Helvetica-Bold", fontSize=12.5, leading=15.5, spaceAfter=3)
    small = ParagraphStyle("s", parent=body, fontSize=7.6, leading=9.8, textColor=colors.HexColor("#444444"))
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=14 * mm,
                            bottomMargin=14 * mm, title=A["title"], author=AUTHOR, creator=ORG,
                            subject="MUAC measurement review, long-rains round 2027", invariant=1)
    s = [Paragraph("%s. To: Kelvin Omollo, %s. From: Winny Karisa, nutrition information officer. 22 July 2027."
                   % (ORG, P.EWU), small),
         Paragraph(A["title"], head)]
    for para in A["paras"]:
        s.append(Paragraph(para, body))
    s += [Image(io.BytesIO(chart_png(A)), width=170 * mm, height=112 * mm), Spacer(1, 2),
          Paragraph(A["footnote"], small)]
    doc.build(s)
    postprocess.scrub(path, P.UNIT, "2027-07-22")


def write(folder, A, test):
    os.makedirs(folder, exist_ok=True)
    _workbook(os.path.join(folder, XLSX), A, test)
    _note(os.path.join(folder, PDF), A)
