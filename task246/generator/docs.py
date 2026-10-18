"""task246 documents: the county's nutrition surveillance protocol (pdf), the June early warning bulletin (pdf), the
bulletin desk's request (eml), the survey teams workbook, PRK's anthropometry and standardisation workbooks (xlsx), the
data dictionary and the readme (md)."""
import datetime as dt
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess

COUNTY = P.COUNTY_GOV
PEOPLE = {"cnio": ("Winny Karisa", "winny.karisa", "laga.go.ke"),
          "supervisor": ("Purity Maina", "purity.maina", "laga.go.ke"),
          "ncha": ("Barasa Kahindi", "b.kahindi", "ncha-kenya.org"),
          "prk": ("Caroline Nthenya", "cnthenya", "plainsfield.or.ke"),
          "ewu": ("Kelvin Omollo", "kelvin.omollo", "laga.go.ke")}


def _addr(k):
    name, user, domain = PEOPLE[k]
    return "%s <%s@%s>" % (name, user, domain)


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.6, leading=13.4, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=13.5, leading=17, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=14,
                        spaceBefore=8, spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10.5, textColor=colors.HexColor("#444444"))
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=8.4, leading=10.2)
    return body, h1, h2, small, cell


def _grid(rows, widths, cell):
    t = Table([[Paragraph(str(c), cell) for c in r] for r in rows], colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA0AB")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E3E8EF"))]))
    return t


def _footer(text):
    def foot(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(20 * mm, 10 * mm, text)
        canvas.drawRightString(A4[0] - 20 * mm, 10 * mm, "Page %d" % d.page)
        canvas.restoreState()
    return foot


def _doc(path, title, author, subject):
    return SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                             bottomMargin=18 * mm, title=title, author=author, creator=author, subject=subject,
                             invariant=1)


# ------------------------------------------------------------------ the request (eml)
def request(path, F):
    text = f"""From: {_addr("ewu")}
To: {_addr("cnio")}
Cc: {_addr("ncha")}, {_addr("prk")}
Date: Wed, 14 Jul 2027 09:12:40 +0300
Subject: Long-rains round figures for the July bulletin
Message-ID: <20270714091240.5512@laga.go.ke>
MIME-Version: 1.0
Content-Type: text/plain; charset="utf-8"
Content-Transfer-Encoding: 7bit

Winny,

The July bulletin goes to the County Steering Group on 30 July, so I need the long-rains round's figures by Friday
23 July: each zone's acute malnutrition figure by MUAC with its alert band, the measurement note wherever the protocol
calls for it, and the mortality and household response figures for the annex.

Barasa and Caroline have sent their anthropometry exports and standardisation sheets; they are in the round's folder
on the shared drive with the teams workbook and the household form exports from the county server.

Thanks,
Kelvin

--
Kelvin Omollo
Drought Information Officer, {P.EWU}
{COUNTY}
"""
    with open(path, "w", newline="\n") as fh:
        fh.write(text)


# ------------------------------------------------------------------ the protocol (pdf)
def protocol(path, F):
    body, h1, h2, small, cell = _styles()
    doc = _doc(path, "Laga County nutrition surveillance: long-rains assessment protocol", "%s, %s" % (P.UNIT, COUNTY),
               "Protocol, version 2.1")
    s = []
    s.append(Paragraph("%s, %s" % (COUNTY, P.DEPT), small))
    s.append(Paragraph("Nutrition surveillance: long-rains assessment protocol", h1))
    s.append(Paragraph("Version 2.1, May 2027. Prepared by the %s; reviewed by the County Nutrition Technical Forum; "
                       "approved by the County Director of Health Services." % P.UNIT, small))
    s.append(Spacer(1, 4))
    sec = [
        ("1. Purpose",
         ["The long-rains assessment round measures acute malnutrition among children aged 6 to 59 months in each of the "
          "county's twelve survey zones (Annex A), and mortality in the zones' households, once a year after the long "
          "rains. The round's figures go to the %s for its monthly bulletin, where each zone's acute malnutrition figure "
          "sets the zone's nutrition alert band (section 8)." % P.EWU]),
        ("2. Sample",
         ["In each zone the round visits 20 clusters (census enumeration areas drawn with probability proportional to "
          "size) and in each cluster 15 households drawn by simple random sampling from the team's listing. Every child "
          "aged 6 to 59 months in an interviewed household is measured.",
          "A household not at home at the first visit is visited again before the team leaves the cluster. A dwelling "
          "found empty, with no household living in it, is recorded as empty and is not visited again. Households are "
          "not replaced."]),
        ("3. Teams",
         ["Twelve survey teams take part, six from each of the two partners, %s (NCHA) and %s (PRK). Each team has a "
          "team leader, an anthropometrist who takes every measurement the team makes, and an assistant who holds the "
          "child and records. The teams workbook lists each team's members and any change of member during the round."
          % (P.PARTNER["A"], P.PARTNER["B"])]),
        ("4. Measurement",
         ["Mid-upper arm circumference (MUAC) is measured on the left arm at the midpoint between the shoulder and the "
          "elbow, the arm hanging relaxed, with the county's standard tapes, and read to the nearest millimetre. "
          "Bilateral pitting oedema is tested by pressing both feet for three seconds."]),
        ("5. Standardisation",
         ["Every anthropometrist and the survey supervisor measure the MUAC of the same ten children twice at the "
          "standardisation session at the end of training (10 June 2027), and of ten other children twice at the "
          "refresher session on the morning of 28 June 2027. An anthropometrist who joins the round after the training "
          "takes the refresher only. The supervisor's mean of her two readings is each test child's reference value: "
          "the supervisor's technique is the round's standard for arm circumference. An anthropometrist's readings "
          "from the sessions they attended are taken together."]),
        ("6. The bulletin's acute malnutrition figure",
         ["For each zone the bulletin publishes the prevalence of acute malnutrition by MUAC among the children aged 6 to "
          "59 months the round measured in the zone, in per cent to one decimal place. A child with bilateral pitting "
          "oedema is counted as malnourished whatever its MUAC. For the other children, because a zone's sample is "
          "small for a count, the share with MUAC below 125 mm is read from the normal distribution with the mean and "
          "the standard deviation of their MUAC.",
          "A child's age is its age in completed months on the day it was measured. A MUAC below 75 mm or above 220 mm "
          "is a recording error, and that child is left out. Each child counts once; a zone's children are those "
          "measured in its clusters."]),
        ("7. The measurement note",
         ["Every MUAC in the round is an anthropometrist's reading of a child. A zone's figure carries the measurement "
          "note when its anthropometrists' error moves it by two percentage points or more, either way: when the figure "
          "read from the zone's readings, as in section 6, and the figure read in the same way from the arm "
          "circumference of the zone's children themselves, at the round's standard (section 5), differ by 2.0 points "
          "or more. The two figures are compared as the bulletin prints them, to one decimal place. A zone with the "
          "note goes into the bulletin at the figure from its children's own arm circumference, marked with the note."]),
        ("8. Alert bands",
         ["A zone's nutrition alert band follows its published figure: Normal below 5.0 per cent, Alert from 5.0 to "
          "9.9, Alarm from 10.0 to 14.9 and Emergency at 15.0 or more."]),
        ("9. Mortality",
         ["A zone's crude death rate is the deaths among its interviewed households' members during the recall period "
          "per 10,000 people per day: its deaths divided by its person-days, times 10,000. A household's person-days "
          "are its mid-point population times its recall days. The mid-point population is the members present at the "
          "interview, less half of those who joined the household or were born into it during the recall period, plus "
          "half of those who left it or died during the recall period. A household's recall days are the date of its "
          "interview less the date of its sub-county's start event (Annex B), in days.",
          "A death is a member who died during the recall period, wherever the death happened, a child born and dying "
          "within the period included. The under-five death rate is worked the same way for children under five: "
          "deaths of children under five over the person-days of the household's children under five, a member's age "
          "being the age the roster gives."]),
        ("10. Household response",
         ["A zone's household response rate is the households interviewed over the households selected, less the "
          "dwellings found empty, in per cent to one decimal place. A household's outcome is the outcome of its last "
          "visit."]),
        ("11. Data",
         ["Each partner enters its teams' anthropometry on its own tablet form and sends its export to the %s: NCHA's "
          "form records MUAC in millimetres and the child's date of birth; PRK's records MUAC in centimetres and the "
          "child's age in completed months. Both partners' teams enter the household interview on the county's "
          "household form, whose server exports the household visits, the member rosters and the births module." % P.UNIT]),
    ]
    for head, paras in sec:
        s.append(Paragraph(head, h2))
        for p in paras:
            s.append(Paragraph(p, body))
    s.append(Paragraph("Annex A. Survey zones", h2))
    rows = [["Zone", "Sub-county", "Clusters"]]
    for i, (z, name, sub, *_r) in enumerate(P.ZONES):
        rows.append([name, sub, "%d to %d" % ((i + 1) * 100 + 1, (i + 1) * 100 + P.CLUSTERS_PER_ZONE)])
    s.append(_grid(rows, [60 * mm, 45 * mm, 35 * mm], cell))
    s.append(Paragraph("Annex B. Recall period start events, 2027 round", h2))
    rows = [["Sub-county", "Start event", "Date"]]
    for sub in ["Laga North", "Laga Central", "Laga East", "Laga West"]:
        d, ev = P.RECALL_START[sub]
        rows.append([sub, ev, d.strftime("%d %B %Y").lstrip("0")])
    s.append(_grid(rows, [45 * mm, 50 * mm, 40 * mm], cell))
    s.append(Paragraph("Document history", h2))
    rows = [["Version", "Date", "Change"],
            ["1.0", "May 2025", "First issue."],
            ["2.0", "May 2026", "Probit reading of the bulletin's figure (section 6); the measurement note (section 7)."],
            ["2.1", "May 2027", "Refresher session added mid-round (section 5); the under-five death rate (section 9)."]]
    s.append(_grid(rows, [20 * mm, 25 * mm, 115 * mm], cell))
    doc.build(s, onFirstPage=_footer("%s: long-rains assessment protocol v2.1" % COUNTY),
              onLaterPages=_footer("%s: long-rains assessment protocol v2.1" % COUNTY))
    postprocess.scrub(path, P.UNIT, "2027-05-27")


# ------------------------------------------------------------------ the June bulletin (pdf, a distractor)
def bulletin(path, F, R):
    body, h1, h2, small, cell = _styles()
    doc = _doc(path, "Laga County drought early warning bulletin, June 2027", P.EWU, "Monthly bulletin")
    s = []
    s.append(Paragraph("%s, %s" % (COUNTY, P.EWU), small))
    s.append(Paragraph("Drought early warning bulletin: June 2027", h1))
    s.append(Paragraph("Issued 8 July 2027. Phase classification by livelihood zone, with the month's rainfall, "
                       "pasture, livestock, market and nutrition indicators.", small))
    s.append(Paragraph("Overall phase", h2))
    s.append(Paragraph("The long rains ended in the third dekad of May. Totals were near normal in Laga North and "
                       "Laga Central and below normal in much of Laga East and Laga West, where pasture is already "
                       "declining and trekking distances to water have lengthened. The county remains at the Alert "
                       "phase, with the pastoral zones of Laga East and Laga West watched for a move to Alarm.", body))
    s.append(Paragraph("Indicators by sub-county", h2))
    rows = [["Sub-county", "Long rains, per cent of normal", "Pasture", "Goat to maize terms of trade (kg)",
             "CHV screening: MUAC below 125 mm"],
            ["Laga North", "96", "Fair", "96", "4.1% of 3,412 screened"],
            ["Laga Central", "91", "Fair", "101", "3.6% of 2,988 screened"],
            ["Laga East", "72", "Poor", "84", "5.2% of 3,105 screened"],
            ["Laga West", "68", "Poor", "79", "5.8% of 2,861 screened"]]
    s.append(_grid(rows, [28 * mm, 30 * mm, 20 * mm, 36 * mm, 40 * mm], cell))
    s.append(Paragraph("Community health volunteers screen children at household visits with colour-banded tapes and "
                       "refer those in the red and yellow bands; the screening shares describe the children they "
                       "reached, not a sample of the zone, and are not comparable with the long-rains round.", small))
    s.append(Paragraph("Phase by livelihood zone", h2))
    rows = [["Livelihood zone", "Phase", "Trend"],
            ["Pastoral (North, Central)", "Alert", "Stable"], ["Pastoral (East, West)", "Alert", "Worsening"],
            ["Agro-pastoral", "Normal", "Stable"], ["Lakeshore and riverine", "Normal", "Stable"],
            ["Laga Town", "Normal", "Stable"]]
    s.append(_grid(rows, [60 * mm, 30 * mm, 30 * mm], cell))
    s.append(Paragraph("Coming next", h2))
    s.append(Paragraph("The long-rains assessment round's fieldwork runs from 14 June to 8 July; its zone figures "
                       "will appear in the July bulletin.", body))
    doc.build(s, onFirstPage=_footer("%s: drought early warning bulletin, June 2027" % COUNTY),
              onLaterPages=_footer("%s: drought early warning bulletin, June 2027" % COUNTY))
    postprocess.scrub(path, P.EWU, "2027-07-08")


# ------------------------------------------------------------------ workbooks
def _sheet_widths(ws, widths):
    for k, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(k)].width = w


def _sheet(wb, title, df, head=None, widths=None, freeze="A2"):
    ws = wb.create_sheet(title)
    r0 = 1
    if head:
        for line in head:
            ws.cell(row=r0, column=1, value=line)
            r0 += 1
        ws.cell(row=1, column=1).font = Font(bold=True)
        r0 += 1
    for j, c in enumerate(df.columns, 1):
        cell = ws.cell(row=r0, column=j, value=c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE3EA")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, row in enumerate(df.itertuples(index=False), r0 + 1):
        for j, v in enumerate(row, 1):
            if v is None or (isinstance(v, float) and v != v):
                continue
            ws.cell(row=i, column=j, value=v.item() if hasattr(v, "item") else v)
    ws.freeze_panes = "A%d" % (r0 + 1) if freeze else None
    if widths:
        _sheet_widths(ws, widths)
    return ws


def teams_workbook(path, R):
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Teams", R.teams(), widths=[9, 6, 18, 18, 14, 34, 34, 11, 11])
    _sheet(wb, "Changes during the round", R.changes(), widths=[11, 9, 6, 100])
    wb.save(path)
    postprocess.scrub(path, P.UNIT, "2027-07-12")
    postprocess.finish_office(path, "2027-06-03T07:40:00Z", "2027-07-12T15:22:00Z", COUNTY)


def prk_anthro_workbook(path, R):
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Children", R.prk_children(), widths=[11, 6, 8, 5, 9, 5, 12, 8, 11, 11, 10, 13])
    wb.save(path)
    postprocess.scrub(path, P.PARTNER["B"], "2027-07-10")
    postprocess.finish_office(path, "2027-06-14T16:05:00Z", "2027-07-10T11:48:00Z", P.PARTNER["B"])


def prk_test_workbook(path, R):
    wb = Workbook()
    wb.remove(wb.active)
    for title, df in R.prk_test_frames():
        _sheet(wb, title, df, head=["%s: MUAC standardisation, %s" % (P.PARTNER["B"], title),
                                    "Readings in cm. Supervisor: %s." % P.SUPERVISOR[0]],
               widths=[9] + [14] * (len(df.columns) - 1))
    wb.save(path)
    postprocess.scrub(path, P.PARTNER["B"], "2027-06-28")
    postprocess.finish_office(path, "2027-06-10T13:30:00Z", "2027-06-28T12:10:00Z", P.PARTNER["B"])


# ------------------------------------------------------------------ the registers
def dictionary(F):
    return f"""# Data dictionary

## {F["sample"]}
The round's cluster sample: one row per cluster. sub_county, zone, cluster (the cluster number), ea_code (the census
enumeration area), households_listed (the team's listing), partner and team (the team that worked it) and
date_visited (dd/mm/yyyy).

## {F["ncha_anthro"]}
NCHA's tablet export, one row per form sent. _uuid and _submission_time (the server's receipt, UTC); today (the day of
measurement, yyyy-mm-dd); team (NCHA team number); cluster; hh_no (the household's number in the cluster's selection);
child_no (the child's line on the household roster); sex (m, f); dob (date of birth, dd/mm/yyyy); oedema (y, n);
weight_kg; height_cm; muac_mm (millimetres); measles (1 card, 2 recall, 3 not vaccinated, 8 does not know; blank for
children under 9 months).

## {F["prk_anthro"]}
PRK's export, sheet Children, one row per child: Date (dd/mm/yyyy); Team (PRK team number); Cluster; HH; Child line;
Sex (1 male, 2 female); Age (months) in completed months; Oedema (1 yes, 0 no); Weight (kg); Height (cm); MUAC (cm);
Measles (Yes, card / Yes, recall / No / Don't know; blank for children under 9 months).

## {F["ncha_test"]}
NCHA's standardisation sheet, one row per reading: session (1 training, 2 refresher); date; test_child; measured_by (the
NCHA anthropometrist's code, or SUPERVISOR); reading (1 or 2); muac_mm.

## {F["prk_test"]}
PRK's standardisation sheets, one per session, one row per test child: the supervisor's two readings, then each PRK
anthropometrist's two readings, in cm.

## {F["teams"]}
Sheet Teams: each team's partner, number, team leader, anthropometrist and the anthropometrist's code, the zones and
clusters it worked, and its first and last fieldwork days. Sheet Changes during the round: dated changes to the teams.

## {F["visits"]}
The county server's household visits (both partners' teams enter the household interview on the county's household
form): one row per visit. partner; team; cluster; hh_no (the household's number in the cluster's selection); visit (1
first, 2 return); date (dd/mm/yyyy); time; outcome (as version 4 of the household form codes it: 1 interviewed, 2
refused, 3 dwelling empty, 4 absent); members_present (interviewed households).

## {F["members"]}
The county server's member rosters: one row per member listed on an interviewed household's roster. cluster; hh_no;
form_id; interview_date (yyyy-mm-dd); line_no; sex (M, F); age and age_unit (completed months for children under five,
completed years otherwise); status (P present at the interview, L left the household during the recall period, D died
during the recall period); joined and born (Yes, No: whether the member joined the household, or was born into it,
during the recall period); left_reason (members who left, as version 4 of the household form codes it: 1 looking for
work, 2 marriage, 3 moved with the mother or caregiver, 4 died away from home, 5 school, 6 moved with the livestock, 7
joined another household); date_of_death (dd/mm/yyyy, members who died at home).

## {F["births"]}
The county server's births module: each child born to the household during the recall period. cluster; hh_no;
form_id; date_of_birth (dd/mm/yyyy); sex; outcome (alive, died); date_of_death.

## {F["bulletin"]}
The {P.EWU}'s June 2027 bulletin.

## {F["markets"]}
Monthly market prices: month (yyyy-mm), market, commodity (with its unit), price (KES).

## {F["rain"]}
Dekadal rainfall at the county's five gauges, January to July 2027: station, dekad (yyyy-mm-D1 to D3), rainfall_mm,
long_term_mean_mm.
"""


def readme(F):
    return f"""# Laga County long-rains assessment round, 2027

Files gathered by the County Nutrition Unit after the 2027 long-rains assessment round, for the July early warning
bulletin.

- {F["request"]}: the bulletin desk's request.
- {F["protocol"]}: the round's protocol (version 2.1).
- {F["sample"]}: the cluster sample.
- {F["teams"]}: the survey teams and changes during the round.
- {F["ncha_anthro"]}: NCHA's anthropometry export.
- {F["prk_anthro"]}: PRK's anthropometry export.
- {F["ncha_test"]}: NCHA's MUAC standardisation sheet.
- {F["prk_test"]}: PRK's MUAC standardisation sheets.
- {F["visits"]}: the household visits, from the county's household form.
- {F["members"]}: the household member rosters, from the county's household form.
- {F["births"]}: the births module, from the county's household form.
- {F["bulletin"]}: the June 2027 early warning bulletin.
- {F["markets"]}: market prices, January 2025 to June 2027.
- {F["rain"]}: dekadal rainfall, 2027.
- DATA_DICTIONARY.md: the fields of each file.
"""


def write_all(target, F, R):
    p = lambda k: os.path.join(target, F[k])
    request(p("request"), F)
    protocol(p("protocol"), F)
    bulletin(p("bulletin"), F, R)
    teams_workbook(p("teams"), R)
    prk_anthro_workbook(p("prk_anthro"), R)
    prk_test_workbook(p("prk_test"), R)
    with open(p("dict"), "w") as fh:
        fh.write(dictionary(F))
    with open(p("readme"), "w") as fh:
        fh.write(readme(F))
