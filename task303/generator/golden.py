"""task303 goldens: the filing memo hardening_memo_2027.pdf and the candidate tables hardening_tables_2027.xlsx, written
from figures.json."""
import io

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess

PDF = "hardening_memo_2027.pdf"
XLSX = "hardening_tables_2027.xlsx"
CAND = ["F-112", "F-114", "F-205", "F-208", "F-301", "F-303"]
HEAD = PatternFill("solid", fgColor="DCE6F0")


def c(n):
    return "{:,}".format(n)


def _chart(fig):
    order = sorted(CAND, key=lambda f: fig["candidates"][f]["exact"])
    vals = [fig["candidates"][f]["exact"] / 1000.0 for f in order]
    labels = ["%s  %s" % (f, fig["candidates"][f]["substation"]) for f in order]
    plt.rcParams["font.family"] = "DejaVu Sans"
    f, ax = plt.subplots(figsize=(6.6, 2.9), dpi=150)
    cols = ["#1F5A8C" if x == fig["call"]["feeder"] else "#A9B8C6" for x in order]
    ax.barh(labels, vals, color=cols, height=0.62)
    for y, f_ in enumerate(order):
        ax.text(vals[y] + 3, y, c(fig["candidates"][f_]["figure"]), va="center", fontsize=8)
    ax.set_xlabel("Customer-minutes removed a year (thousands)", fontsize=8.5)
    ax.tick_params(labelsize=8.5)
    ax.set_xlim(0, max(vals) * 1.22)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.set_title("2027 backbone conversion: customer-minutes each candidate would remove", fontsize=9.5, loc="left")
    f.tight_layout()
    buf = io.BytesIO()
    f.savefig(buf, format="png", metadata={"Software": None})
    plt.close(f)
    buf.seek(0)
    return buf


def memo(path, fig):
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=9.6, leading=13.4, spaceAfter=6)
    head = ParagraphStyle("h", parent=body, fontName="Helvetica-Bold", fontSize=10.2, spaceBefore=6, spaceAfter=3)
    title = ParagraphStyle("t", parent=body, fontName="Helvetica-Bold", fontSize=13, leading=17, spaceAfter=2)
    small = ParagraphStyle("s", parent=body, fontSize=8.4, leading=11, textColor=colors.HexColor("#444444"))
    cl = fig["call"]
    w = cl["feeder"]
    ru = cl["runner_up"]
    sh = fig["shared"]
    story = [Paragraph("%s | Distribution Planning" % P.ORG, small),
             Paragraph("2027 backbone conversion", title),
             Paragraph("Memo for the 2027 Storm Hardening Plan filing, due to the commission 1 December 2026. "
                       "October 2026.", small), Spacer(1, 6),
             Paragraph("Recommendation", head),
             Paragraph("Convert the backbone of <b>%s (%s)</b> in 2027. Valued the way the plan values every project, "
                       "the conversion removes <b>%s customer-minutes</b> of sustained interruption a year, worth "
                       "<b>USD %s a year</b> at the commission's incentive of USD %.2f a customer-minute. The next "
                       "candidate, %s (%s), removes %s." % (w, cl["substation"], c(cl["figure"]), c(cl["value"]),
                                                            P.INCENTIVE, ru, fig["candidates"][ru]["substation"],
                                                            c(cl["runner_up_figure"])), body)]
    hc = ParagraphStyle("hc", parent=body, fontName="Helvetica-Bold", fontSize=8.4, leading=10, spaceAfter=0)
    rows = [[Paragraph(h, hc) for h in ("Feeder", "Substation", "Backbone miles", "Estimated cost (USD million)",
                                        "Customer-minutes removed a year")]]
    for f in sorted(CAND, key=lambda x: fig["candidates"][x]["rank"]):
        k = fig["candidates"][f]
        rows.append([f, k["substation"], "%.1f" % k["backbone_miles"], "%.1f" % k["cost_musd"], c(k["figure"])])
    t = Table(rows, colWidths=[18 * mm, 28 * mm, 24 * mm, 36 * mm, 44 * mm])
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 8.6),
                           ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.6),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F0")),
                           ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#EEF3F8")),
                           ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
                           ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA5B1")),
                           ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    story += [t, Spacer(1, 6), Image(_chart(fig), width=158 * mm, height=66 * mm)]
    pc = [f for f in ("F-112", "F-114") if sh[f]["tickets"]]
    story += [
        Paragraph("Why %s is not first" % fig["default_pick"], head),
        Paragraph("Pell Creek's three feeders (F-112, F-114 and F-117) opened their breakers in the same minute %d "
                  "times outside the major event days from 2019 to 2025. Each time the outage system logged the "
                  "interruption on each feeder's own breaker, with the cause the crew found on the line that supplies "
                  "the substation. A fault on one feeder's backbone does not take its sibling feeders' customers off, "
                  "so these are losses of the substation's supply, and converting a feeder's backbone does not "
                  "prevent them. Credited as faults on the backbone they would put %s at %s customer-minutes a year "
                  "and F-114 at %s; without them %s removes %s and F-114 %s. Byers Mill's feeders share %d such "
                  "minutes; Sorrel Hill's share none." % (sh["F-112"]["tickets"], fig["default_pick"],
                                                         c(fig["default"]["F-112"]["figure"]),
                                                         c(fig["default"]["F-114"]["figure"]), "F-112",
                                                         c(fig["candidates"]["F-112"]["figure"]),
                                                         c(fig["candidates"]["F-114"]["figure"]),
                                                         sh["F-301"]["tickets"]), body),
        Paragraph("How the figures were built", head),
        Paragraph("Both outage exports for 2019 to 2025: the legacy feeder numbers and cause codes mapped to the "
                  "current system's, its UTC stamps put on Eastern time. Tickets that began on a major event day the "
                  "commission accepted are left out. Of the rest, the feeder breaker and line recloser tickets caused "
                  "by vegetation, wind or wildlife, less the tickets every feeder at the substation opened in the same "
                  "minute; customer-minutes are customers interrupted times minutes to restoration; nine in ten of "
                  "them, divided by seven. Each candidate's storm record (the accepted major event days of 2023 to "
                  "2025) and its taps' record are in %s." % XLSX, body)]
    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=16 * mm,
                            bottomMargin=16 * mm, title="2027 backbone conversion", author="Distribution Planning",
                            creator=P.ORG, subject="2027 Storm Hardening Plan", invariant=1)
    doc.build(story)
    postprocess.scrub(path, P.ORG, "2026-10-21")


def tables(path, fig):
    wb = Workbook()
    ws = wb.active
    ws.title = "Conversion 2027"
    ws["A1"] = "2027 backbone conversion: customer-minutes of sustained interruption removed a year"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("2019 to 2025 outage record outside the accepted major event days; feeder breaker and line recloser "
                "tickets from vegetation, wind or wildlife, less those every feeder at the substation opened in the same "
                "minute; nine in ten, divided by seven.")
    ws["A2"].font = Font(italic=True, size=9)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:F2")
    ws.row_dimensions[2].height = 30
    hdr = ["Feeder", "Substation", "Backbone miles", "Estimated cost (USD million)", "Customer-minutes removed a year",
           "Value a year at USD 2.10 (USD)"]
    _head(ws, 4, hdr)
    r = 5
    for f in sorted(CAND, key=lambda x: fig["candidates"][x]["rank"]):
        k = fig["candidates"][f]
        val = fig["call"]["value"] if f == fig["call"]["feeder"] else None
        for j, v in enumerate([f, k["substation"], k["backbone_miles"], k["cost_musd"], k["figure"], val], 1):
            ws.cell(row=r, column=j, value=v)
        ws.cell(row=r, column=5).number_format = "#,##0"
        ws.cell(row=r, column=6).number_format = "#,##0"
        ws.cell(row=r, column=3).number_format = "0.0"
        ws.cell(row=r, column=4).number_format = "0.0"
        r += 1
    ws.cell(row=r + 1, column=1, value="Recommended: %s (%s)." % (fig["call"]["feeder"], fig["call"]["substation"]))
    ws.cell(row=r + 1, column=1).font = Font(bold=True)
    _widths(ws, [10, 14, 10, 14, 18, 18])

    ws = wb.create_sheet("Storm days 2023-25")
    ws["A1"] = "Customers' record on the commission's accepted major event days of 2023 to 2025"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Tickets that began on an accepted major event day (Eastern). Restoration: the crew's Restored status "
                "where the storm crews' records hold the ticket (contract crews' devices on Eastern standard time), "
                "otherwise the outage system's restoration.")
    ws["A2"].font = Font(italic=True, size=9)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:E2")
    ws.row_dimensions[2].height = 30
    _head(ws, 4, ["Feeder", "Substation", "Customer-minutes", "Customers interrupted", "Minutes per customer "
                                                                                       "interrupted"])
    r = 5
    for f in CAND:
        s = fig["storm"][f]
        for j, v in enumerate([f, fig["candidates"][f]["substation"], s["cmi"], s["ci"], s["per"]], 1):
            ws.cell(row=r, column=j, value=v)
        ws.cell(row=r, column=3).number_format = "#,##0"
        ws.cell(row=r, column=4).number_format = "#,##0"
        ws.cell(row=r, column=5).number_format = "0.0"
        r += 1
    _widths(ws, [10, 14, 16, 14, 16])

    ws = wb.create_sheet("Taps 2019-25")
    ws["A1"] = "Interruptions the tap fuses cleared outside the major event days, 2019 to 2025"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Fuse tickets that did not begin on an accepted major event day. Legacy tickets entered from the "
                "trouble crews' handhelds: durations read as hours and minutes, dates as the day the ticket closed.")
    ws["A2"].font = Font(italic=True, size=9)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:D2")
    ws.row_dimensions[2].height = 30
    _head(ws, 4, ["Feeder", "Substation", "Customer-minutes a year", "Minutes per customer interrupted"])
    r = 5
    for f in CAND:
        s = fig["taps"][f]
        for j, v in enumerate([f, fig["candidates"][f]["substation"], s["year"], s["per"]], 1):
            ws.cell(row=r, column=j, value=v)
        ws.cell(row=r, column=3).number_format = "#,##0"
        ws.cell(row=r, column=4).number_format = "0.0"
        r += 1
    _widths(ws, [10, 14, 18, 18])
    wb.properties.creator = "Distribution Planning"
    wb.properties.lastModifiedBy = "Distribution Planning"
    wb.save(path)
    postprocess.scrub(path, P.ORG, "2026-10-21")
    postprocess.finish_office(path, "2026-10-14T13:10:00Z", "2026-10-21T17:32:00Z", P.ORG)


def _head(ws, row, cols):
    for j, h in enumerate(cols, 1):
        x = ws.cell(row=row, column=j, value=h)
        x.font = Font(bold=True)
        x.fill = HEAD
        x.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 32


def _widths(ws, w):
    for j, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(j)].width = x
