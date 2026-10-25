"""task303 documents: the 2027 storm hardening plan's valuation guidance (docx), the outage exports' notes and the
mobile workforce export's notes (txt), the commission's 2024 order on the reliability mechanism (pdf), the feeder
inventory (xlsx), and the two distractors (the vegetation work orders, csv; the line recloser settings, xlsx)."""
import csv
import datetime as dt

import numpy as np
from docx import Document
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

import params as P
import world as W

COMMISSION = "State Utility Regulation Commission"


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Times-Roman", fontSize=10.6, leading=14.6,
                          spaceAfter=6)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Times-Bold", fontSize=12.6, leading=16, spaceAfter=4,
                        alignment=1)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Times-Bold", fontSize=10.8, leading=14,
                        spaceBefore=8, spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=9.0, leading=12, alignment=1)
    return body, h1, h2, small


def _footer(text):
    def foot(canvas, d):
        canvas.saveState()
        canvas.setFont("Times-Roman", 8.4)
        canvas.drawString(22 * mm, 12 * mm, text)
        canvas.drawRightString(LETTER[0] - 22 * mm, 12 * mm, "Page %d" % d.page)
        canvas.restoreState()
    return foot


def order(path):
    body, h1, h2, small = _styles()
    story = [Paragraph(COMMISSION.upper(), h1),
             Paragraph("Docket No. E-2024-031", small),
             Paragraph("In the Matter of the Reliability Performance Mechanism of %s" % P.ORG, small),
             Paragraph("ORDER APPROVING STIPULATION", small),
             Paragraph("Issued: March 14, 2024", small), Spacer(1, 8)]
    paras = [
        ("I. Background", [
            "On September 29, 2023, %s (the Company) filed a proposed reliability performance mechanism under the "
            "Commission's distribution reliability rules. Following discovery and two settlement conferences, the "
            "Company, the Commission's staff and the Office of Consumer Counsel filed a stipulation on February 2, "
            "2024. No party opposed it." % P.ORG]),
        ("II. Reliability reporting", [
            "The Company shall continue to report SAIDI, SAIFI and CAIDI for its system and for each feeder each "
            "calendar year, counting sustained interruptions (longer than five minutes) and excluding major event days. "
            "A major event day is a day on which the Company's daily system SAIDI exceeds the threshold set for that "
            "year by the 2.5 beta method of IEEE Std 1366, computed on the Company's five preceding years. The "
            "Company shall claim its major event days in its annual reliability report, and the Commission shall "
            "accept or reject each claimed day in its annual reliability order."]),
        ("III. The performance incentive", [
            "Beginning with calendar year 2024, the Company shall earn USD 2.10 for each customer-minute of sustained "
            "interruption outside major event days by which its annual record falls below its benchmark, and shall "
            "credit customers USD 2.10 for each customer-minute by which it exceeds the benchmark. The benchmark is "
            "the Company's average for 2019 through 2023, as filed in its 2023 annual reliability report.",
            "The Company's storm hardening plans shall value each project's reduction in customer-minutes at the "
            "same rate."]),
        ("IV. Storm hardening plans", [
            "The Company shall file its storm hardening plan for the following year by December 1 of each year. For "
            "each project the plan shall state its estimated cost, the customer-minutes of sustained interruption a "
            "year the project would remove and their value at the incentive rate, and the reliability record of the "
            "feeder or facility it serves, including its record on the accepted major event days of the three "
            "preceding years."]),
        ("V. Ordering paragraphs", [
            "IT IS THEREFORE ORDERED that the stipulation filed February 2, 2024 is approved, and that the Company "
            "shall file compliance tariff sheets within thirty days of this order.",
            "BY THE COMMISSION."])]
    for head, ps in paras:
        story.append(Paragraph(head, h2))
        for t in ps:
            story.append(Paragraph(t, body))
    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=24 * mm, rightMargin=24 * mm, topMargin=20 * mm,
                            bottomMargin=22 * mm, title="Order Approving Stipulation, Docket No. E-2024-031",
                            author=COMMISSION, creator=COMMISSION, subject="Reliability performance mechanism",
                            invariant=1)
    doc.build(story, onFirstPage=_footer("Docket No. E-2024-031"), onLaterPages=_footer("Docket No. E-2024-031"))


def plan(path):
    F = P.FILES
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    p = d.add_paragraph()
    r = p.add_run("%s | Distribution Planning" % P.ORG)
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    d.add_heading("2027 Storm Hardening Plan: project valuation", level=1)
    p = d.add_paragraph()
    r = p.add_run("Guidance for the plan due to the commission on 1 December 2026. %s, distribution planning, 6 "
                  "October 2026." % P.PLANNER)
    r.italic = True
    r.font.size = Pt(9.5)
    d.add_heading("1. The 2027 backbone conversion", level=2)
    d.add_paragraph("The 2027 budget funds the conversion of one feeder's three-phase backbone, from the feeder "
                    "breaker to the end of the three-phase line, from bare conductor on wood poles to covered "
                    "conductor on steel poles. The conversion prevents nine in ten of the faults on the converted "
                    "backbone caused by vegetation, wind or wildlife. It does nothing for faults on taps and laterals, "
                    "for faults with any other cause, or for faults elsewhere.")
    d.add_heading("2. Candidates", level=2)
    d.add_paragraph("Six feeders in the %s are candidates for 2027. Costs are the planning estimate of USD %.2f "
                    "million a backbone mile." % (P.DISTRICT, P.COST_PER_MILE))
    t = d.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ["Feeder", "Substation", "Backbone miles", "Estimated cost (USD million)"]):
        c.text = h
    for f in W.CAND:
        fp = W.FEEDERS[f]
        row = t.add_row().cells
        row[0].text = f
        row[1].text = P.SUBS[fp[0]][0]
        row[2].text = "%.1f" % fp[2]
        row[3].text = "%.1f" % (fp[2] * P.COST_PER_MILE)
    d.add_heading("3. How a project is valued", level=2)
    d.add_paragraph("Every project in the plan is valued at the customer-minutes of sustained interruption it would "
                    "have removed from the 2019 to 2025 outage record had it been in place, leaving out the major "
                    "event days the commission accepted, divided by seven. Customer-minutes are the customers "
                    "interrupted times the minutes from interruption to restoration. The plan funds the project with "
                    "the largest figure.")
    d.add_heading("4. The incentive", level=2)
    d.add_paragraph("Under the commission's reliability performance mechanism (Docket E-2024-031) the company earns "
                    "USD %.2f for each customer-minute of sustained interruption a year it removes outside major "
                    "event days, and the plan states each project's figure at that rate." % P.INCENTIVE)
    d.add_heading("5. The candidate pages", level=2)
    d.add_paragraph("The filing gives each candidate a page with its feeder's record: its customers' record on the "
                    "major event days of the last three years, and the interruptions on its taps outside the major "
                    "event days, which the conversion leaves as they are.")
    d.add_heading("6. Records", level=2)
    for k, txt in [("legacy", "the outage management system used until 31 December 2022"),
                   ("current", "the outage management system in use from 1 January 2023"),
                   ("meds", "the major event days the commission accepted, 2019 to 2025"),
                   ("inventory", "the district's feeders")]:
        p = d.add_paragraph(style="List Bullet")
        r = p.add_run(F[k])
        r.bold = True
        p.add_run(": " + txt + ".")
    d.add_paragraph("The outage exports' fields are described in %s." % F["notes"])
    d.core_properties.author = P.PLANNER
    d.core_properties.title = "2027 Storm Hardening Plan: project valuation"
    d.save(path)


def notes(path):
    t = """Outage records: export notes (Distribution Operations, September 2026)

Two exports cover 2019 to 2025 for the Valley district's 18 feeders. The company moved to its current outage
management system on 1 January 2023; tickets from before that date are in the legacy export.

oms_interruptions_2019_2022.csv (legacy system)
  TKT_NO    ticket number: two digits of the year, then a number issued in the order the control room opened
            tickets
  FDR       feeder number (112 is feeder F-112)
  SUB_CD    substation code: PLC Pell Creek, SRH Sorrel Hill, BYM Byers Mill, HLF Hollow Fork, KTR Kettle Run,
            WLC Willow Crossing
  OUT_DT    ticket date
  OUT_TM    clock time the interruption began, Eastern local time
  DUR_MIN   duration from the interruption to restoration: whole minutes, or hours and minutes after a point
            (2.15 is 2 hours 15 minutes)
  CUST_OUT  customers interrupted on the ticket
  CAUSE_CD  TREE vegetation, WIND, ANML wildlife, LTNG lightning, EQUP equipment failure, VEH vehicle or public,
            UNK unknown
  DEV_TYP   protective device recorded on the ticket: CB feeder breaker, RCL line recloser, FUS fuse, TX transformer
  DEV_ID    device identifier

oms_interruptions_2023_2025.csv (current system)
  interruption_id, feeder, substation
  interruption_start_utc, restoration_utc   timestamps in UTC; on the days the company ran storm operations the
            storm desk entered restorations from the crews' reports
  customers_interrupted
  cause_category, protective_device_type, protective_device   as above, in the current system's wording

Both exports hold sustained interruptions only (more than five minutes); each ticket is one feeder. The feeder
breaker is the circuit breaker at the substation where the feeder starts; line reclosers sit on the three-phase
backbone; fuses protect the single- and two-phase taps and laterals; a transformer ticket is a distribution
transformer or a single service.

Major event days in major_event_days_2019_2025.csv are Eastern calendar days. An interruption belongs to the day it
began.
"""
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(t)


def crew_notes(path):
    t = """Storm crew status records: export notes (Mobile Workforce, September 2026)

storm_crew_status_2023_2025.csv holds every status the Valley district's storm jobs went through in the mobile
workforce system on the days the company ran storm operations from 2023 to 2025, all of them major event days. One
row per status change, oldest first.

  job_id               mobile workforce job
  oms_interruption_id  the outage ticket the job was opened for; blank for jobs with no outage ticket
  job_type             Outage restoration, Wire down, Damage assessment
  feeder, device       the feeder, and the protective device on the outage ticket
  crew_id, crew_type   the crew and its kind: company line crew, contract line crew, mutual assistance crew,
                       damage assessor
  status               Assigned (the storm desk gave the crew the job), En route, On site, Restored (the crew
                       reported the customers on the job back in supply), Released (the crew left the job)
  status_time_local    the time recorded with the status, to the minute

A job a first crew made safe and left for a second crew carries both crews' statuses.
"""
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(t)


def _xlsx(path, title, header, rows, widths, fmts=None):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCE6F0")
        c.alignment = Alignment(vertical="top", wrap_text=True)
    for r in rows:
        ws.append(list(r))
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for col, f in (fmts or {}).items():
        j = header.index(col) + 1
        for k in range(2, len(rows) + 2):
            ws.cell(row=k, column=j).number_format = f
    ws.freeze_panes = "A2"
    wb.save(path)


def inventory(path, taps):
    rng = np.random.default_rng(P.SEED_DOCS)
    header = ["Feeder", "Substation", "Customers (31 Dec 2025)", "Three-phase backbone (miles)",
              "Taps and laterals (miles)", "Tap fuses", "Line reclosers", "Backbone canopy share (LiDAR 2024)",
              "Poles", "Median pole age (years)", "Backbone conductor"]
    rows = []
    for f, fp in W.FEEDERS.items():
        s, cust, bb, lat, can, zones, fac, cand = fp
        rows.append([f, P.SUBS[s][0], cust, bb, lat, len(taps[f]), sum(1 for z in zones[1:] if z is not None),
                     round(can, 2), int(bb * 26 + lat * 22 + rng.integers(-40, 40)), int(rng.integers(24, 47)),
                     ["336 ACSR", "4/0 ACSR", "1/0 ACSR"][int(rng.integers(0, 3))]])
    _xlsx(path, "Feeders", header, rows, [9, 15, 13, 14, 13, 9, 9, 15, 8, 12, 13],
          {"Customers (31 Dec 2025)": "#,##0", "Backbone canopy share (LiDAR 2024)": "0%"})


def recloser_settings(path):
    rng = np.random.default_rng(P.SEED_DOCS + 1)
    header = ["Recloser", "Feeder", "Location (pole)", "Control", "Phase minimum trip (A)", "Ground minimum trip (A)",
              "Phase curves", "Ground curves", "Operations to lockout", "Reclose intervals (s)", "Last tested"]
    rows = []
    for f, fp in W.FEEDERS.items():
        zones = fp[5]
        for z in (1, 2):
            if zones[z] is None:
                continue
            n = f[2:]
            pick = int(rng.integers(0, 2))
            rows.append(["F%s-R%d" % (n, z), f, "%s-%04d" % (n, int(rng.integers(300, 1900)) + 400 * z),
                         ["Electronic, type E3", "Electronic, type E4"][pick],
                         [400, 280][z - 1] - 40 * int(rng.integers(0, 3)), [140, 100][z - 1],
                         "A fast, C delay", "A fast, D delay" if pick else "1 fast, 1 delay", 4,
                         ["2 - 10 - 15", "1 - 5 - 15"][int(rng.integers(0, 2))],
                         (dt.date(2024, 3, 1) + dt.timedelta(days=int(rng.integers(0, 700)))).isoformat()])
    _xlsx(path, "Line reclosers", header, rows, [10, 8, 12, 18, 12, 12, 14, 14, 11, 14, 11])


def vegetation(path):
    rng = np.random.default_rng(P.SEED_DOCS + 2)
    rows = []
    k = 0
    crews = ["Crew A", "Crew B", "Crew C"]
    for f, fp in W.FEEDERS.items():
        secs = 4
        first = int(rng.integers(0, 4))
        for sidx in range(secs):
            yr = 2017 + (first + sidx) % 4
            while yr <= 2026:
                k += 1
                day = dt.date(yr, int(rng.integers(1, 12)), int(rng.integers(1, 28)))
                if day <= dt.date(2026, 9, 15):
                    rows.append(["VM-%05d" % k, f, "%s-S%d" % (f, sidx + 1), round((fp[2] + fp[3]) / secs, 1),
                                 day.isoformat(), crews[int(rng.integers(0, 3))]])
                yr += 4
    rows.sort(key=lambda r: (r[4], r[0]))
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["work_order", "feeder", "section", "section_miles", "completed", "contractor_crew"])
        w.writerows(rows)
