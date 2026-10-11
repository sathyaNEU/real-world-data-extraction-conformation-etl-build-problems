"""The documents in the pack: methodology note (PDF), the SEC data-set readme (HTM), the client
watchlist (XLSX), the client flag-file layout (MD), the data-sources note (MD), and two files the
solution does not use (short interest CSV, the SEC 13(f) list TXT)."""
import csv
import datetime as dt
import os

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

import world as W
import postprocess as PP

FIRM = "Quillfield Research LLC"
METHOD_PDF = "ownership_methodology_OM-4.pdf"
README = "FORM13F_readme.htm"
WATCHLIST = "client_watchlist_crowding_2022Q4.xlsx"
LAYOUT = "client_flag_file_layout_v3.md"
SOURCES = "data_sources.md"
SHORTS = "short_interest_20221230.csv"
LIST13F = "13flist2022q4.txt"


def write_docs(wd, target):
    method_pdf(os.path.join(target, METHOD_PDF))
    readme(os.path.join(target, README))
    watchlist(wd, os.path.join(target, WATCHLIST))
    layout(os.path.join(target, LAYOUT))
    sources(wd, os.path.join(target, SOURCES))
    shorts(wd, os.path.join(target, SHORTS))
    list13f(wd, os.path.join(target, LIST13F))


# ------------------------------------------------------------------ methodology note
METHOD = [
    ("h", "Ownership methodology note OM-4"),
    ("s", f"{FIRM} &middot; Methodology &middot; revision 4, 18 October 2022 &middot; owner: Nancy Henderson"),
    ("p", "This note fixes how the ownership desk measures institutional ownership from Form 13F and how the "
          "crowded flag in client flag files is set. It replaces OM-3 (March 2021). Changes from OM-3: section 5 "
          "now names the cover-page tag; section 6 covers the move of 13F values to whole dollars."),
    ("b", "1. Sources"),
    ("p", "Holdings come from Form 13F reports as filed on EDGAR: the SEC's Form 13F data sets for closed "
          "submission quarters, and the data engineering team's nightly build from EDGAR information-table XML for "
          "the current quarter, in the same layout. Share counts come from the cover pages of issuers' 10-K and "
          "10-Q filings, as carried in the financial statement data sets."),
    ("b", "2. Institutional ownership"),
    ("p", "Institutional ownership of an issuer for a report period is the number of its shares held by all "
          "managers that file a Form 13F holdings report or combination report for that period, each report taken "
          "as amended, divided by the issuer's shares outstanding. Managers that file a 13F notice report no "
          "holdings of their own; their holdings sit in another manager's report."),
    ("b", "3. What counts as a holding"),
    ("p", "Information-table lines reported in shares (SH) with no put or call entry. Put and call lines are "
          "exposure, not ownership, and are left out. A manager's holding of a security is the sum of its lines "
          "for that security, whatever the investment discretion or other-manager entries on them."),
    ("b", "4. Matching lines to securities"),
    ("p", "Lines are matched to the issuer's common stock on CUSIP. Filers key CUSIPs themselves: a CUSIP is read "
          "as nine characters, upper case, with leading zeros restored. A line whose CUSIP fails its check digit is "
          "matched on issuer name and title of class where exactly one covered security fits."),
    ("b", "5. Shares outstanding"),
    ("p", "The issuer's cover-page count (dei:EntityCommonStockSharesOutstanding) with the latest as-of date on or "
          "before the report period end."),
    ("b", "6. Market value and holders"),
    ("p", "Market value is the information-table value on a whole-dollar basis. Holders are named at the level of "
          "the filing manager (one CIK); affiliates that file their own reports are not combined."),
    ("b", "7. Crowded flag"),
    ("p", "An issuer carries the crowded flag for a report period when its institutional ownership for that "
          "period is above 80.0 per cent."),
    ("b", "8. Client flag files"),
    ("p", "One flag file per report period and client watchlist, issued on the date in the client calendar and not "
          "reissued. The file layout is client flag file layout v3."),
    ("s", "Internal. Questions to the methodology mailbox. Next scheduled review: April 2023."),
]


def method_pdf(path):
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["BodyText"], fontName="Helvetica", fontSize=9.5, leading=12.5, spaceAfter=5)
    head = ParagraphStyle("h", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=14, spaceAfter=4)
    sub = ParagraphStyle("s", parent=body, fontSize=8, textColor="#555555", spaceAfter=10)
    sec = ParagraphStyle("c", parent=body, fontName="Helvetica-Bold", fontSize=10, spaceBefore=4, spaceAfter=2)
    flow = []
    for kind, text in METHOD:
        flow.append(Paragraph(text, {"h": head, "s": sub, "p": body, "b": sec}[kind]))
    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=0.9 * inch, rightMargin=0.9 * inch,
                            topMargin=0.8 * inch, bottomMargin=0.8 * inch, title="OM-4 Ownership methodology",
                            author="Nancy Henderson", subject="Methodology", creator="Microsoft Word",
                            producer="Microsoft: Print To PDF")
    doc.build(flow)
    PP.normalize_pdf(path, stamp=b"D:20221018142233Z'", doc_id=b"9C2E51B07A4F3D18E6B2A0C94D7F1E35")


# ------------------------------------------------------------------ SEC data-set readme
def readme(path):
    rows = [
        ("SUBMISSION", [
            ("ACCESSION_NUMBER", "The 20-character string formed from the 18-digit number assigned by the Commission to each EDGAR submission."),
            ("FILING_DATE", "Date the submission was filed (DD-MON-YYYY)."),
            ("SUBMISSIONTYPE", "Type of submission: 13F-HR, 13F-HR/A, 13F-NT or 13F-NT/A."),
            ("CIK", "Central Index Key of the filing manager."),
            ("PERIODOFREPORT", "Period of report (DD-MON-YYYY)."),
        ]),
        ("COVERPAGE", [
            ("ACCESSION_NUMBER", "Accession number of the submission."),
            ("REPORTCALENDARORQUARTER", "Report for the calendar year or quarter ended (DD-MON-YYYY)."),
            ("ISAMENDMENT", "Y if the submission is an amendment."),
            ("AMENDMENTNO", "Amendment number."),
            ("AMENDMENTTYPE", "RESTATEMENT or NEW HOLDINGS. A restatement amends the report in its entirety; a new holdings amendment adds holdings entries to the report it amends."),
            ("CONFDENIEDEXPIRED", "Y if the amendment reports holdings for which confidential treatment was denied or has expired."),
            ("FILINGMANAGER_NAME", "Name of the institutional investment manager filing the report."),
            ("REPORTTYPE", "13F HOLDINGS REPORT, 13F NOTICE or 13F COMBINATION REPORT."),
            ("FORM13FFILENUMBER", "Form 13F file number assigned to the filing manager."),
        ]),
        ("INFOTABLE", [
            ("ACCESSION_NUMBER", "Accession number of the submission."),
            ("INFOTABLE_SK", "Surrogate key for the information-table entry."),
            ("NAMEOFISSUER", "Name of issuer, as reported."),
            ("TITLEOFCLASS", "Title of class, as reported."),
            ("CUSIP", "CUSIP number of the security, as reported."),
            ("VALUE", "Market value of the holding. For filings made before January 3, 2023, rounded to the nearest one thousand dollars (x$1000); for filings made on or after January 3, 2023, rounded to the nearest dollar."),
            ("SSHPRNAMT", "Shares or principal amount."),
            ("SSHPRNAMTTYPE", "SH for shares, PRN for principal amount."),
            ("PUTCALL", "Put or Call, where the line reports an option."),
            ("INVESTMENTDISCRETION", "SOLE, DFND (shared-defined) or OTR (shared-other)."),
            ("OTHERMANAGER", "Sequence numbers of other included managers, from the list of other included managers."),
            ("VOTING_AUTH_SOLE / SHARED / NONE", "Voting authority, in shares."),
        ]),
    ]
    h = ["<!DOCTYPE html><html><head><meta charset=\"utf-8\"><title>Form 13F Data Sets - Readme</title>",
         "<style>body{font-family:Arial,Helvetica,sans-serif;font-size:13px;max-width:980px;margin:20px}"
         "table{border-collapse:collapse;margin-bottom:18px}td,th{border:1px solid #999;padding:3px 6px;vertical-align:top}"
         "th{background:#e8e8e8}</style></head><body>",
         "<h2>Form 13F Data Sets</h2>",
         "<p>The Form 13F data sets below provide information extracted from Form 13F filings submitted to the "
         "Commission. Each data set contains the filings submitted during the stated date range, whatever the "
         "period the filing reports. Institutional investment managers that use the United States mail and exercise "
         "investment discretion over $100 million or more in Section 13(f) securities file Form 13F within 45 days "
         "after the end of each calendar quarter.</p>",
         "<p>The data is presented as filed. It has not been corrected or adjusted, and the Commission does not "
         "warrant its accuracy. Users should consult the filings themselves where precision matters.</p>",
         "<p>Tables are tab-delimited, UTF-8, with a header row. Tables included in this extract: SUBMISSION, "
         "COVERPAGE, INFOTABLE. Joins are on ACCESSION_NUMBER.</p>"]
    for t, fields in rows:
        h.append(f"<h3>{t}</h3><table><tr><th>Field</th><th>Description</th></tr>")
        for f_, d_ in fields:
            h.append(f"<tr><td>{f_}</td><td>{d_}</td></tr>")
        h.append("</table>")
    h.append("<p style=\"font-size:11px;color:#555\">Division of Economic and Risk Analysis. Last updated January 2023.</p></body></html>")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(h) + "\n")


# ------------------------------------------------------------------ client watchlist
def watchlist(wd, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Watchlist"
    ws.append(["Client", "Harrowby Point Partners"])
    ws.append(["Watchlist", "Crowding monitor, US mid cap"])
    ws.append(["Effective", "2022-10-03"])
    ws.append([])
    hdr = ["Seq", "Ticker", "Issuer", "CUSIP", "Issuer CIK", "Share class", "Added"]
    ws.append(hdr)
    for c in ws[5]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE5F0")
    added = ["2021-04-01", "2021-04-01", "2022-01-03", "2021-04-01", "2021-10-01", "2022-07-01",
             "2021-04-01", "2022-04-01", "2021-07-01", "2022-10-03", "2021-04-01", "2022-07-01"]
    for k, i in enumerate([x for x in wd.iss if x["watch"]]):
        ws.append([k + 1, i["ticker"], _disp(i["name"]),
                   i["cusip"], i["issuer_cik"], "Common", added[k]])
    for col, w in zip("ABCDEFG", [6, 9, 38, 13, 12, 12, 12]):
        ws.column_dimensions[col].width = w
    ws2 = wb.create_sheet("Notes")
    ws2.append(["Watchlist agreed with the client's risk team. Changes take effect from the next report period."])
    ws2.append(["Ticker changes are tracked in the security master; CUSIP is the key."])
    wb.properties.creator = "Client Services"
    wb.properties.lastModifiedBy = "Client Services"
    wb.properties.created = dt.datetime(2021, 3, 29, 15, 2, 11)
    wb.properties.modified = dt.datetime(2022, 10, 3, 13, 12, 44)
    wb.save(path)
    PP.normalize_zip(path, modified="2022-10-03T13:12:44Z")


def _disp(name):
    words = name.title().split()
    if words[-1] in ("Inc", "Corp", "Co", "Ltd"):
        words[-1] += "."
    return " ".join(words)


# ------------------------------------------------------------------ client flag file layout
def layout(path):
    txt = """# Client flag file layout, v3

Owner: client services. In use since the 2021Q2 files.

One CSV per report period and watchlist, UTF-8, comma-separated, header row, one row per watchlist
security in watchlist sequence.

| Column | Content |
|---|---|
| `report_period` | Report period end, YYYY-MM-DD |
| `cusip` | Nine-character CUSIP from the watchlist |
| `ticker` | Ticker from the watchlist |
| `shares_outstanding` | The share count institutional ownership is measured against (OM-4 section 5), whole shares |
| `inst_own_pct` | Institutional ownership for the report period, per cent |
| `crowded` | `Y` or `N` |
| `prior_shares_outstanding` | As `shares_outstanding`, for the previous report period |
| `prior_inst_own_pct` | As `inst_own_pct`, for the previous report period |
| `top_holder_cik` | CIK of the largest holder by market value at the report period end |
| `top_holder_name` | That manager's name as on its 13F cover page |
| `top_holder_value_usd` | That holder's market value, whole US dollars |

Rows for securities with no institutional holders carry `0.0` and `N` and leave the holder columns
blank.
"""
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(txt)


# ------------------------------------------------------------------ data sources
def sources(wd, path):
    txt = f"""# Data sources for the ownership desk

Maintained by Jonathan Richardson (data engineering). Pull for the December 2022 quarter files.

| Path | What it is | Coverage | How it is produced |
|---|---|---|---|
| `form13f_2022q3/` | SEC Form 13F data set, tables SUBMISSION, COVERPAGE, INFOTABLE | filings submitted 1 Jul to 30 Sep 2022 | downloaded from sec.gov as published, unaltered |
| `form13f_2022q4/` | SEC Form 13F data set, same tables | filings submitted 1 Oct to 31 Dec 2022 | downloaded from sec.gov as published, unaltered |
| `edgar13f_daily_20230103_20230203/` | nightly build from EDGAR 13F information-table XML, in the data-set layout | filings accepted 3 Jan to 3 Feb 2023 | parsed by our loader; INFOTABLE_SK is our own surrogate key |
| `fsds_cover_extract/` | cover-page share counts in the financial statement data sets layout (sub, num) | 10-Q and 10-K filings of covered issuers accepted by 3 Feb 2023 | the security master's pull; num carries the dei cover tags only |
| `manager_register.csv` | the desk's register of 13F filing managers | every manager we have covered since 2009 | maintained by hand; `family` groups affiliates for research notes, not for ownership |
| `{WATCHLIST}` | the client's watchlist | 12 securities | from client services |
| `{LAYOUT}` | the client flag file layout | | client services |
| `{METHOD_PDF}` | ownership methodology | | methodology |
| `{README}` | SEC readme for the Form 13F data sets | | sec.gov |
| `{SHORTS}` | exchange short interest for the watchlist | settlement 30 Dec 2022 | vendor feed |
| `{LIST13F}` | SEC Official List of Section 13(f) Securities, fourth quarter 2022, rows for covered issuers | | sec.gov, trimmed to covered CUSIPs |

The SEC data sets are organised by the quarter a filing was submitted in, not the period it
reports. Data-set and daily-build files are kept as delivered; nothing in them has been corrected.
"""
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(txt)


# ------------------------------------------------------------------ distractors
def shorts(wd, path):
    r = W.rng("shorts")
    rows = []
    for i in [x for x in wd.iss if x["watch"]]:
        si = int(i["S_oct"] * r.uniform(0.018, 0.11))
        prev = int(si * r.uniform(0.85, 1.15))
        adv = int(i["S_oct"] * r.uniform(0.004, 0.012))
        rows.append({"settlementDate": "2022-12-30", "symbolCode": i["ticker"],
                     "issueName": i["name"].title() + " Common Stock", "marketClassCode": "NNM" if r.random() < 0.5 else "NYSE",
                     "currentShortPositionQuantity": si, "previousShortPositionQuantity": prev,
                     "averageDailyVolumeQuantity": adv, "daysToCoverQuantity": round(si / adv, 2),
                     "changePercent": round(100 * (si - prev) / prev, 2)})
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)


def list13f(wd, path):
    r = W.rng("13flist")
    lines = ["                                      List of Section 13F Securities",
             "                                          Fourth Quarter FY 2022",
             "",
             "Run Date: 11/30/2022                                                   Run Time: 14:02:31",
             "",
             "CUSIP NO      ISSUER NAME                     ISSUER DESCRIPTION       STATUS",
             ""]
    for i in sorted(wd.iss, key=lambda x: x["cusip"]):
        c = i["cusip"]
        star = "*" if r.random() < 0.7 else " "
        lines.append(f"{c[:6]} {c[6:8]} {c[8]}{star}  {i['name'][:30]:<30}  COM                      ")
    lines.append("")
    lines.append(f"Total Count: {len(wd.iss)}")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
