"""the golden deliverables: the root-cause memo (docx, with its chart) and the workings workbook (xlsx)"""
import io
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import estim as E
import postprocess

DOCX = "tree_outage_root_cause_2025.docx"
XLSX = "tree_outage_workings_2025.xlsx"
COMPANY = "Vessey County Electric Cooperative"
GREEN = "2F4F3A"


def f1(x):
    return "{:,.1f}".format(x)


def chart(A, path):
    import numpy as np
    bb = A["by_band"]["North"]
    base = bb.loc[list(E.BASE)].mean()
    now = bb.loc[E.YEAR]
    names = ["Under 1", "1 to 2", "2 to 3", "3 to 4"]
    x = np.arange(4)
    fig, ax = plt.subplots(figsize=(6.4, 3.3), dpi=160)
    ax.bar(x - 0.19, [base[i] for i in range(4)], 0.36, color="#9DB3A3", label="2021 to 2023 average")
    ax.bar(x + 0.19, [now[i] for i in range(4)], 0.36, color="#2F4F3A", label="2025")
    ax.set_xticks(x)
    ax.set_xticklabels(names, fontsize=8)
    ax.set_xlabel("Years since the unit was last trimmed", fontsize=8)
    ax.set_ylabel("Tree-caused outages per 100 miles", fontsize=8)
    ax.tick_params(axis="y", labelsize=8)
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(frameon=False, fontsize=7.5, loc="upper left")
    ax.set_title("North: 2025 rates rose at every time since trim, units trimmed within the year included", fontsize=9,
                 loc="left")
    fig.tight_layout()
    fig.savefig(path, format="png", metadata={"Software": None})
    plt.close(fig)


def _shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)
    pr.append(sh)


def _cell(cell, text, bold=False, white=False, right=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.font.size = Pt(9.5)
    r.bold = bold
    if white:
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    if right:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _footer(d, text):
    p = d.sections[0].footer.paragraphs[0]
    r = p.add_run(text + "    Page ")
    r.font.size = Pt(8)
    for kind, val in (("begin", None), (None, "PAGE"), ("end", None)):
        run = p.add_run()
        run.font.size = Pt(8)
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            run._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = val
            run._r.append(it)


def _source(d, text):
    p = d.add_paragraph(text)
    p.runs[0].font.size = Pt(8)
    p.runs[0].italic = True


def memo(A, path):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for name in ("Heading 1", "Heading 2"):
        d.styles[name].font.name = "Calibri"
        d.styles[name].font.color.rgb = RGBColor(0x2F, 0x4F, 0x3A)
    d.add_heading("2025 tree outages: ash dying beyond the right-of-way, not the trimming", level=1)
    for k, v in (("To", "Lori Rodriguez, Manager of Right-of-Way and Vegetation Management"),
                 ("From", "Reliability Engineering"), ("Date", "14 October 2026"),
                 ("Subject", "Review of the 2025 rise in tree-caused outages (ENG-REL-02 4.1) and the 2027 right-of-way "
                             "budget increase")):
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(k + ": ")
        r.bold = True
        r.font.size = Pt(9)
        r = p.add_run(v)
        r.font.size = Pt(9)
    d.add_paragraph()
    p = d.add_paragraph()
    r = p.add_run("Fund hazard-tree removal beyond the right-of-way with the 2027 budget increase. Do not re-bid the "
                  "trimming contract, and do not put the money into catch-up trimming.")
    r.bold = True
    c = A["comp"]
    d.add_paragraph("The cooperative's tree-caused outage rate was %s outages per 100 miles of overhead primary in 2025, "
                    "against %s on average over 2021 to 2023. North carries most of the rise: its rate went from %s in "
                    "2023 to %s in 2024 and %s in 2025, after its trimming passed to Garrity Vegetation Services on 1 "
                    "January 2024." % (f1(c["rate_2025"]), f1(c["rate_base"]), f1(A["rates"][2023]["North"]),
                                       f1(A["rates"][2024]["North"]), f1(A["rates"][2025]["North"])))
    t = d.add_table(rows=1, cols=6)
    t.style = "Table Grid"
    hdr = ["District", "2021", "2022", "2023", "2024", "2025"]
    for i, h in enumerate(hdr):
        _cell(t.rows[0].cells[i], h, bold=True, white=True, right=i > 0)
        _shade(t.rows[0].cells[i], GREEN)
    for dist in list(E.DISTRICTS) + ["all"]:
        row = t.add_row().cells
        tot = dist == "all"
        _cell(row[0], "Cooperative" if tot else dist, bold=tot)
        for i, y in enumerate(range(2021, 2026)):
            _cell(row[i + 1], f1(A["rates"][y][dist]), bold=tot, right=True)
            if tot:
                _shade(row[i + 1], "E4ECE6")
        if tot:
            _shade(row[0], "E4ECE6")
    _source(d, "Tree-caused outages per 100 miles of overhead primary in service, outside major event days (ENG-REL-02 "
               "2.2). Source: outage management system export; trim-unit register.")
    d.add_heading("Why it is not the contractor", level=2)
    s = A["std"]
    d.add_paragraph(
        "The rise sits on North's units whatever crew trimmed them last. Set against North's 2025 rates at the same "
        "time since trim, the units Garrity last trimmed failed at %s times the rate expected of them, and the units "
        "Hollenbeck last trimmed at %s times. A crew leaving clearance short would show on its own units first; "
        "neither firm's units stand apart." % ("{:.2f}".format(s[E.INCOMING][2]), "{:.2f}".format(s[E.OUTGOING][2])))
    d.add_heading("Why trimming sooner would not stop it", level=2)
    young = A["young"]["North"]
    d.add_paragraph(
        "Read on the regrowth clock, North's 2025 rate rose by about %s outages per 100 miles on units trimmed within "
        "the previous two years, where regrowth seldom reaches the line, about as much as on units near the end of "
        "their cycle. A tree falling in from beyond the right-of-way does that; regrowth inside it does not. The rise "
        "also climbs with the share of ash among the trees the forester's survey counts within falling distance of "
        "each unit: little on North's units with the least ash, most on those with the most, a gradient that was not "
        "there in 2021 to 2023 and appeared in 2024 with the rise. Ash dying beyond the right-of-way and falling onto "
        "the line does that." % f1(young))
    png = io.BytesIO()
    chart(A, png)
    png.seek(0)
    d.add_picture(png, width=Inches(6.0))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _source(d, "Time since trim runs from each unit's last accepted trim, or its in-service date where the line was "
               "rebuilt (OPS-ROW-01 3.1, 3.2). Source: outage management system export; trimming completions.")
    d.add_heading("The other districts", level=2)
    d.add_paragraph(
        "East's 2025 rate (%s) sits on units its April 2025 handover left past the four-year cycle, and South's (%s) "
        "on a windier year; together they carry well under half of North's rise. Catch-up trimming in East answers its "
        "part; it is not where the cooperative's rise came from." % (f1(A["rates"][2025]["East"]),
                                                                      f1(A["rates"][2025]["South"])))
    p = d.add_paragraph("Rates exclude the major event days of %s (IEEE Std 1366, 2.5 beta, from the five prior years). "
                        "Workings are in %s." % (" and ".join(x.strftime("%-d %B") for x in A["med_days"]), XLSX))
    p.runs[0].font.size = Pt(8.5)
    _footer(d, "Vessey County Electric Cooperative, Reliability Engineering. Tree-caused outage review, 2025.")
    d.core_properties.title = "2025 tree outage review"
    d.core_properties.author = "Reliability Engineering"
    d.save(path)
    postprocess.scrub(path, "Reliability Engineering", "2026-10-14")
    postprocess.finish_office(path, "2026-10-08T14:10:00Z", "2026-10-14T16:20:00Z", COMPANY)


HEAD = Font(name="Calibri", bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor=GREEN)


def _sheet(ws, title, header, rows, widths, formats, note):
    last = get_column_letter(len(header))
    ws["A1"] = title
    ws["A1"].font = Font(name="Calibri", bold=True, size=12)
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A1:%s1" % last)
    ws.row_dimensions[1].height = 15.0 * (1 + len(title) // int(sum(widths) * 1.0))
    for i, h in enumerate(header, start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = HEAD
        c.fill = FILL
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left" if i == 1 else "right")
    for r_i, r in enumerate(rows, start=4):
        for c_i, v in enumerate(r, start=1):
            c = ws.cell(row=r_i, column=c_i, value=v)
            if c_i in formats:
                c.number_format = formats[c_i]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    n = 5 + len(rows)
    c = ws.cell(row=n, column=1, value=note)
    c.font = Font(name="Calibri", italic=True, size=9)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A%d:%s%d" % (n, last, n))
    ws.row_dimensions[n].height = 12.5 * (1 + len(note) // int(sum(widths) * 1.15))
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = "A1:%s%d" % (last, n)


def workbook(A, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Contractor billing"
    rows = []
    for v in (E.OUTGOING, E.INCOMING):
        rows.append([v] + [round(A["billing"][(v, q)], 1) for q in ("Q1", "Q2", "Q3", "Q4")])
    _sheet(ws, "Miles of right-of-way billed for trimming done in each quarter of 2025",
           ["Contractor", "Q1", "Q2", "Q3", "Q4"], rows, [36, 11, 11, 11, 11], {2: "#,##0.0", 3: "#,##0.0", 4: "#,##0.0",
                                                                                   5: "#,##0.0"},
           "Trimming lines only, by the month of the work; credit memos taken against the invoices they reverse; "
           "linear feet converted at 5,280 to the mile.")
    ws = wb.create_sheet("Major event days")
    rows = [[dist] + list(A["med"][dist]) for dist in E.DISTRICTS]
    _sheet(ws, "Taken out of the 2025 reliability figures: the major event days of %s" %
           " and ".join(x.strftime("%-d %B 2025") for x in A["med_days"]),
           ["District", "Outages", "Members interrupted", "Customer-minutes"], rows, [14, 12, 20, 18],
           {2: "#,##0", 3: "#,##0", 4: "#,##0"},
           "Unplanned sustained outages on primary and secondary, one per device interruption; customer-minutes over "
           "the restoration steps (ENG-REL-02 3.1 to 3.3).")
    ws = wb.create_sheet("Miles in service")
    rows = [[dist, round(A["miles"][dist], 1)] for dist in E.DISTRICTS]
    _sheet(ws, "Overhead primary in service on 31 December 2025", ["District", "Miles"], rows, [14, 12],
           {2: "#,##0.0"}, "Trim-unit register: units in service on 31 December 2025, retirements dated after it kept.")
    ws = wb.create_sheet("Tree outage rate")
    rows = [["Cooperative" if dist == "all" else dist] + [round(A["rates"][y][dist], 1) for y in range(2021, 2026)]
            for dist in list(E.DISTRICTS) + ["all"]]
    _sheet(ws, "Tree-caused outages per 100 miles of overhead primary in service (ENG-REL-02 2.2)",
           ["District", "2021", "2022", "2023", "2024", "2025"], rows, [14] + [10] * 5,
           {i: "0.0" for i in range(2, 7)}, "2021 to 2023 average for the cooperative: %s; 2025: %s." %
           (f1(A["comp"]["rate_base"]), f1(A["comp"]["rate_2025"])))
    wb.properties.creator = "Reliability Engineering"
    wb.properties.lastModifiedBy = "Reliability Engineering"
    wb.save(path)
    postprocess.scrub(path, "Reliability Engineering", "2026-10-14")
    postprocess.finish_office(path, "2026-10-06T13:00:00Z", "2026-10-14T15:55:00Z", COMPANY)


def write(folder, A):
    memo(A, os.path.join(folder, DOCX))
    workbook(A, os.path.join(folder, XLSX))
