"""task129 generator: the two deliverables and generator/figures.json, written from the check record.

q4_squad_call.pptx carries the call on its first slide and the five changes' figures as a chart on the second;
q4_squad_call_workings.xlsx carries the dated changes by title and the image delivery by month. Office packages
are rewritten with fixed entry stamps and core dates so two builds are byte-identical."""
import io
import json
import os
import re
import zipfile

import params as P

DECK = "q4_squad_call.pptx"
BOOK = "q4_squad_call_workings.xlsx"
AUTHOR = "Elin Lauritsen"
WHEN = "2026-09-11T08:40:00Z"
NAME = {"A": "Header-bidding ad stack", "B": "Image pipeline", "C": "Bølge front end", "D": "Consent banner",
        "E": "Brand web fonts"}
MONTHS = ["2026-%02d" % m for m in range(3, 9)]
MONTH_NAME = {"2026-03": "March", "2026-04": "April", "2026-05": "May", "2026-06": "June", "2026-07": "July",
              "2026-08": "August"}


def r10k(x):
    return int(round(x / 10_000.0)) * 10_000


def r1(x):
    return round(x + 1e-12, 1)


def r10(x):
    return int(round(x / 10.0)) * 10


def answers(rec):
    f = rec["figures"]
    five = {"A": f["A"], "B": f["B"], "C": f["C"], "D": f["D"], "E": f["E"]}
    order = sorted(five, key=lambda k: -five[k])
    return dict(
        call=order[0], runner=order[1], five={k: r10k(v) for k, v in five.items()}, order=order,
        T={x: {t: r10k(f["T"][x][t]) for t in P.TITLES} for x in "DBE"},
        I={m: (r1(rec["I"][m][0]), r10(rec["I"][m][1])) for m in MONTHS},
        default={"A": r10k(f["R1"]["A"]), "C": r10k(f["R1"]["C"])},
    )


def _normalise(path, title):
    with zipfile.ZipFile(path) as z:
        items = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    out = []
    for n, b in items:
        if n == "docProps/core.xml":
            b = re.sub(rb"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\d(\.\d+)?Z", WHEN.encode(), b)
            b = re.sub(rb"<dc:creator>[^<]*</dc:creator>", b"<dc:creator>" + AUTHOR.encode() + b"</dc:creator>", b)
            b = re.sub(rb"<cp:lastModifiedBy>[^<]*</cp:lastModifiedBy>",
                       b"<cp:lastModifiedBy>" + AUTHOR.encode() + b"</cp:lastModifiedBy>", b)
        if n.endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(b)) as zz:
                inner = [(i.filename, zz.read(i.filename)) for i in zz.infolist()]
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zo:
                for m, bb in inner:
                    if m == "docProps/core.xml":
                        bb = re.sub(rb"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\d(\.\d+)?Z", WHEN.encode(), bb)
                    zi = zipfile.ZipInfo(m, date_time=(2026, 9, 11, 8, 40, 0))
                    zi.compress_type = zipfile.ZIP_DEFLATED
                    zo.writestr(zi, bb)
            b = buf.getvalue()
        out.append((n, b))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in out:
            zi = zipfile.ZipInfo(n, date_time=(2026, 9, 11, 8, 40, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            z.writestr(zi, b)


def _footer(slide, n):
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    fb = slide.shapes.add_textbox(Inches(0.6), Inches(6.95), Inches(10.8), Inches(0.35))
    fb.text_frame.text = ("Digital operations, quarterly service review of 17 September 2026. Field data: RUM beacon "
                          "export, 1 February to 31 August 2026; release train log; Bølge rollout log.")
    fb.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    fb.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    nb = slide.shapes.add_textbox(Inches(12.3), Inches(6.95), Inches(0.5), Inches(0.35))
    nb.text_frame.text = str(n)
    nb.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    nb.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x77, 0x77, 0x77)


def deck(A, path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.7), Inches(12.0), Inches(5.8))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.text = "Q4 squad: fix the %s" % NAME[A["call"]]
    tf.paragraphs[0].runs[0].font.size = Pt(34)
    tf.paragraphs[0].runs[0].font.bold = True
    p = tf.add_paragraph()
    p.text = ("Fixing it brings back about {:,} mobile views a month under the line at August's traffic. The closest "
              "other change, the {}, would bring back about {:,}.").format(
        A["five"][A["call"]], NAME[A["runner"]].lower().replace("bølge", "Bølge"), A["five"][A["runner"]])
    p.runs[0].font.size = Pt(20)
    p.space_before = Pt(14)
    w = tf.add_paragraph()
    w.text = ("Read by phone class alone, the ad stack comes first. The front end costs a view mostly on a device's "
              "first view after a release-train deploy and the ad stack mostly on its later views, and first views "
              "are three quarters of the ad-supported views but about a quarter of the puzzles and subscriber views "
              "the two are measured on, so each is carried within phone class and first or later view.")
    w.runs[0].font.size = Pt(15)
    w.space_before = Pt(18)
    q = tf.add_paragraph()
    q.text = ("Each change valued on the scored mobile views it acts on at August's traffic (service level v3, "
              "section 4). The front end and the ad stack switched on together, template group by template group; "
              "the puzzles pages carry the ad stack alone from 24 June and subscribers' ad-free views the front end "
              "alone from 30 June.")
    q.runs[0].font.size = Pt(13)
    q.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    q.space_before = Pt(18)
    _footer(s, 1)
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    t2 = s2.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.0), Inches(0.8))
    t2.text_frame.text = "Fixing the front end brings back more views than any other change"
    t2.text_frame.paragraphs[0].runs[0].font.size = Pt(22)
    t2.text_frame.paragraphs[0].runs[0].font.bold = True
    u2 = t2.text_frame.add_paragraph()
    u2.text = "Mobile views a month back under the line at August 2026 traffic, nearest ten thousand"
    u2.runs[0].font.size = Pt(13)
    u2.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    cd = CategoryChartData()
    cd.categories = [NAME[k] for k in reversed(A["order"])]
    cd.add_series("Mobile views a month", [A["five"][k] for k in reversed(A["order"])])
    gf = s2.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.45), Inches(12.0), Inches(5.3), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = False
    pl = ch.plots[0]
    pl.gap_width = 60
    pl.has_data_labels = True
    pl.data_labels.number_format = "#,##0"
    pl.data_labels.number_format_is_linked = False
    pl.data_labels.font.size = Pt(11)
    pl.series[0].format.fill.solid()
    pl.series[0].format.fill.fore_color.rgb = RGBColor(0x9A, 0xA9, 0xB8)
    top = pl.series[0].points[len(A["order"]) - 1]
    top.format.fill.solid()
    top.format.fill.fore_color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    ch.category_axis.tick_labels.font.size = Pt(12)
    ch.value_axis.tick_labels.font.size = Pt(10)
    ch.value_axis.tick_labels.number_format = "#,##0"
    ch.value_axis.tick_labels.number_format_is_linked = False
    ch.value_axis.has_major_gridlines = False
    _footer(s2, 2)
    prs.core_properties.author = AUTHOR
    prs.core_properties.last_modified_by = AUTHOR
    prs.core_properties.title = "Q4 squad call"
    prs.core_properties.revision = 2
    prs.save(path)
    _normalise(path, "Q4 squad call")


def workbook(A, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    thin = Side(style="thin", color="B7BEC8")
    head = PatternFill("solid", fgColor="2F5D8A")
    ws = wb.active
    ws.title = "Dated changes by title"
    ws["A1"] = "What fixing each dated change would bring back, by title"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Mobile views a month back under the line at August 2026 traffic, nearest 10,000. Each step read over "
                "the 28 days either side of its release, phones only, each page view once, titles as today.")
    ws["A2"].font = Font(italic=True, size=9, color="555555")
    cols = ["Change"] + [P.TITLE_NAME[t] for t in P.TITLES] + ["All three titles"]
    for j, h in enumerate(cols, start=1):
        c = ws.cell(row=4, column=j, value=h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, x in enumerate(["D", "B", "E"]):
        ws.cell(row=5 + i, column=1, value=NAME[x]).border = Border(bottom=thin)
        for j, t in enumerate(P.TITLES, start=2):
            c = ws.cell(row=5 + i, column=j, value=A["T"][x][t])
            c.number_format = "#,##0"
            c.border = Border(bottom=thin)
        c = ws.cell(row=5 + i, column=5, value=A["five"][x])
        c.number_format = "#,##0"
        c.font = Font(bold=True)
        c.border = Border(bottom=thin)
    ws["A9"] = "The all-titles figure is rounded from the unrounded sum of the three titles."
    ws["A9"].font = Font(italic=True, size=9, color="555555")
    ws.column_dimensions["A"].width = 26
    for col in "BCDE":
        ws.column_dimensions[col].width = 18
    ws.freeze_panes = "B5"
    ws.print_title_rows = "4:4"
    w2 = wb.create_sheet("Image delivery")
    w2["A1"] = "Images the CDN served to mobile page views, per mobile page view"
    w2["A1"].font = Font(bold=True, size=12)
    w2["A2"] = ("Phones, browser and app-webview requests (edge hits and origin fills), kilobytes of 1,000 bytes, over "
                "the month's mobile page views.")
    w2["A2"].font = Font(italic=True, size=9, color="555555")
    for j, h in enumerate(["Month", "Image requests per mobile page view", "Image kilobytes per mobile page view"],
                          start=1):
        c = w2.cell(row=4, column=j, value=h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, m in enumerate(MONTHS):
        rq, kb = A["I"][m]
        w2.cell(row=5 + i, column=1, value=MONTH_NAME[m] + " 2026")
        a = w2.cell(row=5 + i, column=2, value=rq)
        a.number_format = "0.0"
        b = w2.cell(row=5 + i, column=3, value=kb)
        b.number_format = "#,##0"
    w2.column_dimensions["A"].width = 16
    w2.column_dimensions["B"].width = 22
    w2.column_dimensions["C"].width = 22
    w2.freeze_panes = "A5"
    w3 = wb.create_sheet("Notes")
    notes = [
        "Notes",
        "",
        "Field data: RUM beacon export, 1 February to 31 August 2026 (Copenhagen dates), exported 5 September 2026; "
        "image CDN delivery log for the same months.",
        "A mobile view is a page view on a device the registry lists as a phone (service level v3, section 1); each "
        "page view is counted once by its page-view key and weighted by sample_weight.",
        "A scored view misses the line when its largest contentful paint exceeds 4.0 seconds; views with no paint "
        "timing are not scored (section 2).",
        "Each change's step is the change in the title's over-line share over the 28 days either side of the moment "
        "its release was deployed (section 3), valued on August 2026's scored views the change now reaches "
        "(section 4).",
        "Titles are reported by current masthead (edition register), as in the Q3 close-out.",
        "Beacons repeating an earlier beacon of the same device for the same page with the same timings within ten "
        "minutes, under another page-view key, are left out.",
        "Image delivery counts smartphone requests from browsers and the app's webview, edge hits and origin fills; "
        "the app's native image clients make no page view and are left out. The fallback vendor reported bytes in "
        "kilobytes from 9 June to 21 July; they are converted. A kilobyte is 1,000 bytes.",
    ]
    for i, line in enumerate(notes, start=1):
        w3.cell(row=i, column=1, value=line).alignment = Alignment(wrap_text=True, vertical="top")
    w3["A1"].font = Font(bold=True, size=12)
    w3.column_dimensions["A"].width = 110
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.save(path)
    _normalise(path, "Q4 squad call workings")


def figures(rec, A, path):
    f = rec["figures"]
    fig = {
        "call": NAME[A["call"]], "call_views_month": "{:,}".format(A["five"][A["call"]]),
        "runner_up": NAME[A["runner"]], "runner_up_views_month": "{:,}".format(A["five"][A["runner"]]),
        "five": {NAME[k]: "{:,}".format(v) for k, v in A["five"].items()},
        "default_path": {"A": "{:,}".format(A["default"]["A"]), "C": "{:,}".format(A["default"]["C"])},
        "T": {NAME[x]: {P.TITLE_NAME[t]: "{:,}".format(A["T"][x][t]) for t in P.TITLES} for x in "DBE"},
        "I": {MONTH_NAME[m]: {"requests_per_view": "%.1f" % A["I"][m][0], "kb_per_view": "{:,}".format(A["I"][m][1])}
              for m in MONTHS},
        "first_view_share_pct": {k: "%d" % round(100 * v) for k, v in rec["first_share"].items()},
        "joint_points": "%.2f" % f["joint_pts"], "raw_points": "%.2f" % f["raw_pts"],
        "puzzles_step_views": "{:,}".format(r10k(f["puz_step"])),
    }
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(fig, fh, ensure_ascii=False, indent=1, sort_keys=True)
        fh.write("\n")
    return fig


def write(rec, golden_dir, fig_path):
    os.makedirs(golden_dir, exist_ok=True)
    A = answers(rec)
    deck(A, os.path.join(golden_dir, DECK))
    workbook(A, os.path.join(golden_dir, BOOK))
    return A, figures(rec, A, fig_path)
