"""task129 documents: the reliability reporting policy (docx), the right-of-way maintenance policy (pdf), the
trimming contracts register (xlsx) and the data dictionary (xlsx)"""
import os

from docx import Document
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess

COMPANY = "Vessey County Electric Cooperative"


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=13, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=14, leading=18, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=14, spaceBefore=8,
                        spaceAfter=4)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10, textColor=colors.HexColor("#444444"))
    return dict(body=body, h1=h1, h2=h2, small=small)


def _footer(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(0.75 * inch, 0.5 * inch, text)
        canvas.drawRightString(letter[0] - 0.75 * inch, 0.5 * inch, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _pdf(path, story, title, author, footer, stamp):
    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.8 * inch, rightMargin=0.8 * inch, topMargin=0.7 * inch,
                            bottomMargin=0.8 * inch, title=title, author=author, creator=author, invariant=1)
    doc.build(story, onFirstPage=_footer(footer), onLaterPages=_footer(footer))
    postprocess.scrub(path, author, stamp)


HEAD = Font(name="Calibri", bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="2F4F3A")


def _sheet(ws, header, rows, widths=None, wrap=False, start=1):
    for i, h in enumerate(header, start=1):
        c = ws.cell(row=start, column=i, value=h)
        c.font = HEAD
        c.fill = FILL
        c.alignment = Alignment(vertical="top", wrap_text=True)
    for r_i, r in enumerate(rows, start=start + 1):
        for c_i, v in enumerate(r, start=1):
            c = ws.cell(row=r_i, column=c_i, value=v)
            c.font = Font(name="Calibri", size=11)
            if wrap:
                c.alignment = Alignment(vertical="top", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=start + 1, column=1)


def _save(wb, path, who, created, modified):
    wb.properties.creator = who
    wb.properties.lastModifiedBy = who
    wb.save(path)
    postprocess.scrub(path, who, modified[:10])
    postprocess.finish_office(path, created, modified, COMPANY)


# ------------------------------------------------------------------------------------------------ reliability policy (docx)
REL = [
    ("1 Purpose", [
        "1.1 This policy sets how the cooperative counts its outages and reports its reliability to the board of "
        "directors and in the annual reliability report to the members. It applies to the distribution system from the "
        "substation breakers to the members' meters."]),
    ("2 Outages and the tree-caused outage rate", [
        "2.1 An outage is an interruption of supply by one device opening, with every member it took out and every "
        "restoration step that brought them back. A device that has opened stays open until it is restored, so a device "
        "is in one outage at a time; the outage management system records outages under ticket numbers. An outage is "
        "sustained when power stays off for more than five "
        "minutes, from the outage's start to its last restoration step; an outage restored within five minutes is "
        "momentary. An outage belongs to the day, and to the year, in which it starts, in Central local time.",
        "2.2 The tree-caused outage rate, for a district or for the cooperative in a year, is the number of sustained "
        "outages on overhead primary with cause code TREE, outside major event days, per 100 miles of overhead primary "
        "in service, the miles averaged over the days of the year from each trim unit's in-service and retirement "
        "dates."]),
    ("3 Reliability indices and major event days", [
        "3.1 The reliability figures (SAIFI, SAIDI and CAIDI) count the unplanned sustained outages on primary and "
        "secondary alike; planned outages are not in them. The members interrupted by an outage are the members it "
        "took out; its customer-minutes add, for each restoration step, the members restored at that step times the "
        "minutes from the outage's start to that step.",
        "3.2 Major event days follow IEEE Std 1366. For each year, the natural logarithms of the daily SAIDI of the "
        "five years before it, leaving out days with no customer-minutes, give a mean and a standard deviation, and a "
        "day of the year whose SAIDI exceeds the exponential of the mean plus 2.5 standard deviations is a major event "
        "day. Daily SAIDI is the customer-minutes of the day's unplanned sustained outages over the members served "
        "at the end of that year.",
        "3.3 Outages on major event days are reported separately and are taken out of the reliability figures and of "
        "the tree-caused outage rate."]),
    ("4 Reliability reviews", [
        "4.1 When the cooperative's tree-caused outage rate in a year stands above its average over the three years "
        "before the year preceding it, Engineering and Operations review the rise. The review names the driver that "
        "accounts for most of the difference and the programme the right-of-way maintenance policy assigns to that "
        "driver, and the programme is carried into the next budget the board approves."]),
    ("5 Records", [
        "5.1 The outage management system is the record of outages. Members served are counted by Member Services at "
        "each year's end. Overhead primary miles are carried in the trim-unit register kept with the engineering GIS."]),
]


def reliability_policy(path):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for name in ("Heading 1", "Heading 2", "Title"):
        d.styles[name].font.name = "Calibri"
        d.styles[name].font.color.rgb = RGBColor(0x2F, 0x4F, 0x3A)
    d.add_heading("Reliability Reporting Policy", level=1)
    p = d.add_paragraph("Vessey County Electric Cooperative. Policy ENG-REL-02, revision 4, adopted by the board of "
                        "directors 21 March 2023. Policy owner: Dennis Garcia, Reliability Engineer.")
    p.runs[0].font.size = Pt(9)
    for head, paras in REL:
        d.add_heading(head, level=2)
        for t in paras:
            d.add_paragraph(t)
    d.core_properties.title = "Reliability Reporting Policy ENG-REL-02"
    d.core_properties.author = "Dennis Garcia"
    d.core_properties.revision = 9
    d.save(path)
    postprocess.scrub(path, "Dennis Garcia", "2023-03-21")
    postprocess.finish_office(path, "2023-02-06T15:02:00Z", "2023-03-21T19:44:00Z", COMPANY)


# ------------------------------------------------------------------------------------------------ right-of-way policy (pdf)
ROW = [
    ("1 Scope", ["This policy covers the right-of-way of the cooperative's overhead primary lines in the North, East and "
                 "South districts. Secondary lines and service drops are cleared on member request."]),
    ("2 Clearance", ["Trimming clears the right-of-way: fifteen feet either side of the pole line, ground to sky. Every "
                     "tree inside it is cut back or removed and brush is cut or treated. Trees standing beyond the "
                     "right-of-way are not trimmed."]),
    ("3 Cycle", ["3.1 Each trim unit is trimmed every four years. A unit is past its cycle when more than four years have "
                 "passed since its last accepted trim.",
                 "3.2 A unit built or rebuilt is cleared at construction, and its cycle runs from its in-service date."]),
    ("4 Contract work", ["4.1 Trimming is let by district to contractors under the contracts in the trimming contracts "
                         "register. Work orders are issued by circuit, and a circuit's units are trimmed as one campaign.",
                         "4.2 The forester audits each completed unit against the clearance in section 2. A unit that "
                         "fails the audit is re-trimmed at the contractor's cost, and its trim is accepted when it passes.",
                         "4.3 Contractors bill the units they complete under their contract's terms, by the month of "
                         "the work."]),
    ("5 Hazard trees", ["A hazard tree is a tree standing beyond the right-of-way that would reach the line if it fell. "
                        "The forester's span survey counts the trees within falling distance of each span by species. "
                        "Hazard-tree removal fells such trees with the landowner's consent."]),
    ("6 Programmes", ["6.1 Catch-up trimming answers units past their cycle.",
                      "6.2 A re-bid of a district's trimming contract answers trimming that falls short of the clearance "
                      "in section 2.",
                      "6.3 Hazard-tree removal answers trees falling onto the line from beyond the right-of-way.",
                      "6.4 Weather below major event days calls for nothing new."]),
]


def row_policy(path):
    s = _styles()
    st = [Paragraph("Right-of-Way Maintenance Policy", s["h1"]),
          Paragraph("Vessey County Electric Cooperative, Operations. Policy OPS-ROW-01, revision 6, effective 2 January "
                    "2024. Policy owner: Lori Rodriguez, Manager of Right-of-Way and Vegetation Management. Approved by "
                    "Brian Montgomery, Manager of Operations.", s["small"]), Spacer(1, 6)]
    for head, paras in ROW:
        st.append(Paragraph(head, s["h2"]))
        for t in paras:
            st.append(Paragraph(t, s["body"]))
    _pdf(path, st, "Right-of-Way Maintenance Policy OPS-ROW-01", "Lori Rodriguez",
         "OPS-ROW-01 rev. 6, effective 2 January 2024", "2024-01-02")


# ------------------------------------------------------------------------------------------------ contracts (xlsx)
def contracts(path):
    rows = [
        ["North", "ROW-N-2020", P.OUTGOING, "2020-01-01", "2023-12-31",
         "Work orders issued before 1 January 2024 completed under this contract."],
        ["North", "ROW-N-2024", P.INCOMING, "2024-01-01", "2028-12-31", ""],
        ["East", "ROW-E-2021", P.OUTGOING, "2021-04-01", "2025-03-31",
         "No new work orders issued after 31 December 2024."],
        ["East", "ROW-E-2025", P.INCOMING, "2025-04-01", "2030-03-31", ""],
        ["South", "ROW-S-2022", P.OUTGOING, "2022-01-01", "2026-12-31", ""],
        ["North", "ROW-N-2016", P.OUTGOING, "2016-01-01", "2019-12-31", ""],
        ["East", "ROW-E-2017", P.OUTGOING, "2017-04-01", "2021-03-31", ""],
        ["South", "ROW-S-2018", P.OUTGOING, "2018-01-01", "2021-12-31", ""],
    ]
    rows.sort(key=lambda r: (r[0], r[3]))
    wb = Workbook()
    ws = wb.active
    ws.title = "Contracts"
    _sheet(ws, ["District", "Contract", "Contractor", "Start", "End", "Notes"], rows, widths=[10, 13, 34, 12, 12, 60],
           wrap=True)
    _save(wb, path, "Lori Rodriguez", "2016-06-14T13:20:00Z", "2025-03-27T16:05:00Z")


# ------------------------------------------------------------------------------------------------ data dictionary (xlsx)
DICT = {
    "outages": ("One row per restoration step of an outage ticket; a ticket's rows share its outage_id.", [
        ("outage_id", "Ticket number in the outage management system"),
        ("start_local", "When the outage began, Central local time"),
        ("level", "PRIMARY (overhead primary line) or SECONDARY (transformer or service drop)"),
        ("service_area", "District"), ("circuit", "Circuit (substation and feeder)"),
        ("unit_id", "Trim unit the opened device or the transformer sits on"),
        ("device_type", "FUSE, RECLOSER, BREAKER, SWITCH, TRANSFORMER or SERVICE"),
        ("device_id", "The device or equipment that opened"),
        ("cause_code", "TREE, ANIMAL, EQUIPMENT, LIGHTNING, VEHICLE, UNKNOWN or PLANNED"),
        ("members_out", "Members the outage took out"), ("step", "Restoration step number"),
        ("step_restored_local", "When power returned to the members restored at this step, Central local time"),
        ("members_restored", "Members restored at this step")]),
    "register": ("One row per trim unit.", [
        ("unit_id", "Trim unit"), ("district", "District"), ("substation", "Substation"), ("circuit", "Circuit"),
        ("miles", "Overhead primary in the unit, miles"), ("in_service_date", "Date the unit entered service"),
        ("retired_date", "Date the unit leaves service"), ("retire_reason", "REBUILT or UNDERGROUND"),
        ("replaced_by", "Unit that replaced a rebuilt unit")]),
    "completions": ("One row per trim of a unit.", [
        ("record_id", "Completion record"), ("work_order", "Work order (one circuit campaign)"),
        ("unit_id", "Trim unit"), ("contractor", "Contractor that did the work"), ("crew", "Crew"),
        ("completed_date", "Date the crew completed the unit"), ("audit_status", "ACCEPTED or REJECTED at the forester's audit; PENDING until the audit is done"),
        ("accepted_date", "Date the trim was accepted")]),
    "survey": ("One row per span and species group with trees found.", [
        ("span_id", "Span (circuit and the pole it starts from)"), ("unit_id", "Trim unit"),
        ("survey_date", "Date surveyed"), ("species_code", "Species group as recorded by the survey tool"),
        ("trees_within_falling_distance", "Trees that would reach the line if they fell")]),
    "weather": ("One row per station and day.", [
        ("station_id", "Station"), ("station_name", "Station name"), ("date_utc", "Day, UTC"),
        ("peak_gust_mph", "Peak gust, miles per hour"), ("precipitation_in", "Precipitation, inches")]),
    "members": ("One row per district and year.", [
        ("district", "District"), ("year", "Year"), ("members_at_year_end", "Members served at 31 December")]),
    "invoices": ("One row per invoice line.", [
        ("vendor", "Contractor"), ("invoice_no", "Invoice or credit memo number"), ("invoice_date", "Document date"),
        ("doc_type", "INVOICE or CREDIT MEMO"), ("applies_to_invoice", "Invoice a credit memo reverses"),
        ("work_order", "Work order billed"), ("service_month", "Month of the work billed"),
        ("line_no", "Line"), ("item_code", "Contract item"), ("description", "Item description"),
        ("quantity", "Quantity billed"), ("uom", "Unit of measure"), ("unit_price", "Unit price, USD"),
        ("amount", "Line amount, USD")]),
}
DICT_FILES = {"outages": "oms_outage_export_2016-2025.csv", "register": "trim_unit_register.csv",
              "completions": "row_trim_completions_2012-2025.csv", "survey": "row_tree_survey_2023-2024.parquet",
              "weather": "weather_daily_2016-2025.csv", "members": "member_counts_2015-2025.csv",
              "invoices": "ap_invoices_row_trimming_2024-12_2025-12.csv"}


def dictionary(path):
    wb = Workbook()
    first = True
    for name, (grain, fields) in DICT.items():
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        ws["A1"] = DICT_FILES[name] + ": " + grain
        ws["A1"].font = Font(name="Calibri", italic=True)
        _sheet(ws, ["Field", "Description"], [list(f) for f in fields], widths=[30, 90], start=3)
    _save(wb, path, "Dennis Garcia", "2021-02-01T14:00:00Z", "2026-01-13T17:30:00Z")


def write_all(target, FILES, U, T, M):
    reliability_policy(os.path.join(target, FILES["rel_policy"]))
    row_policy(os.path.join(target, FILES["row_policy"]))
    contracts(os.path.join(target, FILES["contracts"]))
    dictionary(os.path.join(target, FILES["dictionary"]))
    return {"counts": {}}
