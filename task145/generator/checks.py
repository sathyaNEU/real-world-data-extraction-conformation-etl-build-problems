"""task145 checks, run on every build: the input gates, the three-rules block (R1 to R3c) on the clean truth and the
pack, every rung by name with its placement, every construction of the golden, every semantic fork row (SF1 to SF20),
the bins and flip conditions, the asks with every stop and the cracker's path, the separation line, the distractors,
the realism sweep, the load-bearing facts and the golden files. Any failure stops the build."""
from __future__ import annotations

import builtins
import csv
import datetime as dt
import json
import math
import os
import re
import shutil
import tempfile
import zipfile
from collections import Counter, defaultdict

import openpyxl

import calc as C
import common as K
import field as FD
import records as R
import verify as V

GW = K.GATEWAYS
DASH = re.compile("[%s%s]" % (chr(0x2013), chr(0x2014)))
# a sentence on how the goods are priced (the 2027 plan's "Carrier pricing" is Larchwood's own freight rate increase)
PRICE_WORDS = (r"\bprices?\b|(?<!Carrier )\bpricing\b|\binvoiced at\b|pass.?through|repric|local.currency|"
               r"(?:Canadian|CAD).dollar (?:price|list)")
MAIN_FILES = list(R.F_HV.values()) + [R.F_GW_NEW, R.F_GW_LEGACY, R.F_PLAN27, R.F_STD, R.F_CLR, R.F_NOTE_CLR]
ASK_FILES = [R.F_TA_OLD, R.F_TA_NEW, R.F_HR_STAFF, R.F_HR_CHG, R.F_CC, R.F_NOTE_PAY, R.F_BCS_PDF, R.F_BCS_PORTAL,
             R.F_AP, R.F_BCS_RATES]


class Checks:
    def __init__(self):
        self.res = []

    def __call__(self, label, ok, detail=""):
        self.res.append((label, bool(ok), detail))
        return ok

    def report(self):
        bad = [r for r in self.res if not r[1]]
        for lab, ok, det in self.res:
            if not ok:
                print("FAIL %s %s" % (lab, det))
        print("checks.py: %d assertions, %d failed" % (len(self.res), len(bad)))
        if bad:
            raise SystemExit(1)


class Reads:
    """Which pack files a computation opens."""

    def __init__(self, target):
        self.target = os.path.abspath(target)
        self.seen = set()

    def note(self, p):
        try:
            p = os.path.abspath(p)
            if p.startswith(self.target):
                self.seen.add(os.path.basename(p))
        except TypeError:
            pass

    def __enter__(self):
        self.orig = builtins.open
        self.orig_wb = openpyxl.load_workbook
        me = self

        def o(p, *a, **k):
            me.note(p)
            return me.orig(p, *a, **k)

        def wb(p, *a, **k):
            me.note(p)
            return me.orig_wb(p, *a, **k)
        builtins.open = o
        openpyxl.load_workbook = wb
        return self

    def __exit__(self, *a):
        builtins.open = self.orig
        openpyxl.load_workbook = self.orig_wb


def copy_without(target, drop):
    d = tempfile.mkdtemp(suffix=".t145")
    for f in os.listdir(target):
        if f not in drop:
            shutil.copy(os.path.join(target, f), os.path.join(d, f))
    return d


def edge(x, step):
    """Distance from x to the nearest rounding edge of a bin of width step."""
    q = x / step - 0.5
    return abs(q - round(q)) * step if abs(q - round(q)) <= 0.5 else (1 - abs(q - round(q))) * step


def tup(placed, held):
    return tuple(placed[g] for g in GW) + (held,)


def fmt(placed, held):
    return ", ".join("%s %d" % (K.DESK[g], placed[g]) for g in GW if placed[g]) + "; held %d" % held


def texts_of(target):
    """Every document's text: PDFs, the notes, and every workbook cell."""
    from pypdf import PdfReader
    out = {}
    for f in sorted(os.listdir(target)):
        p = os.path.join(target, f)
        if f.endswith(".pdf"):
            out[f] = re.sub(r"\s+", " ", " ".join(pg.extract_text() for pg in PdfReader(p).pages))
        elif f.endswith(".txt"):
            out[f] = re.sub(r"\s+", " ", open(p, encoding="utf-8").read())
        elif f.endswith(".docx"):
            out[f] = V.docx_text(p)
        elif f.endswith(".xlsx"):
            wb = openpyxl.load_workbook(p, read_only=True)
            out[f] = " | ".join(str(c) for ws in wb.worksheets for r in ws.iter_rows(values_only=True) for c in r
                                if c is not None)
    return out


def e_of(n, ratio):
    return {g: n[g] * ratio[g] / 12.0 for g in GW}


def place(E, est=None):
    return C.placement(E, est)


def run(target, golden, W=None, F=None, G=None, meta_path=None):
    import build as B
    import deliver as DL
    A = Checks()
    if W is None or F is None:
        W, F = B.world()
    if G is None:
        G = DL.golden(W, F.staff, F.lines, F.stmts)
    files = sorted(os.listdir(target))

    # ------------------------------------------------------------------ the pack's own computation and what it reads
    with Reads(target) as rd:
        c = V.main_call(target)
    main_reads = set(rd.seen)
    with Reads(target) as rd:
        Ap = V.ask_a(target)
    a_reads = set(rd.seen)
    with Reads(target) as rd:
        Bp = V.ask_b(target)
    b_reads = set(rd.seen)

    # ------------------------------------------------------------------ input gates
    fmts = {f.rsplit(".", 1)[1] for f in files}
    rows = {}
    for f in files:
        if f.endswith(".csv"):
            with open(os.path.join(target, f), encoding="utf-8") as fh:
                rows[f] = sum(1 for _ in fh) - 1
    A("gate: 10 or more files", len(files) >= 10, len(files))
    A("gate: 3 or more formats", len(fmts) >= 3, sorted(fmts))
    A("gate: a file of 25,000 rows or more", max(rows.values()) >= 25000, max(rows.items(), key=lambda kv: kv[1]))
    meta = json.load(open(meta_path)) if meta_path and os.path.exists(meta_path) else {}
    dis = meta.get("distractor_files", [])
    A("gate: two or more distractors declared and shipped", len(dis) >= 2 and all(d in files for d in dis), dis)
    A("gate: no file named for the stump, a distractor, a memo or an answer",
      not [f for f in files if re.search(r"distract|decoy|known.?issue|answer|trap|memo|draft|brief|readme|summary", f,
                                         re.I)], files)

    # ------------------------------------------------------------------ R1: the field's raw material ships
    A("R1 the call reads the four high-value check extracts, the gateway report, the 2027 plan, BRK-STD-03 and "
      "CLR-OPS-11", set(V.MAIN) <= main_reads, sorted(main_reads))
    A("R1 the call reads no ask file and no distractor", not (main_reads & (set(ASK_FILES) | set(dis))),
      sorted(main_reads & (set(ASK_FILES) | set(dis))))
    reg = c["reg"]
    reg_all = c["reg_all"]
    A("R1 the register holds every base-year shipment of the check once (%d)" % len(reg), len(reg) == len(W.ships),
      (len(reg), len(W.ships)))
    A("R1 the eight extracts hold the two years on file, each shipment once (%d and %d)" % (len(W.prior), len(W.ships)),
      len(reg_all) == len(W.prior) + len(W.ships), (len(reg_all), len(W.prior) + len(W.ships)))
    recs = C.records(W)
    pi_w = C.golden_shares(W, recs)
    n_w = C.mixture_counts(recs, pi_w)
    A("R1 the decisive computation reads each shipment's importer, tariff item, currency, invoice value and CAD value in "
      "both years: the pack's pass-through shares and 2027 counts equal the truth's",
      {k: v[2] for k, v in c["pt"].items()} == pi_w and all(abs(c["n"][g] - n_w[g]) < 1e-6 for g in GW),
      {g: round(c["n"][g], 2) for g in GW})
    A("R1 every shipment re-routed at the plan's rates: the pack's counts equal the truth's",
      c["nb"] == {g: C.formal_at(W, g, lambda s: K.PLAN_RATE[s.cur]) for g in GW}, c["nb"])
    heads = []
    for f in files:
        p = os.path.join(target, f)
        if f.endswith(".csv"):
            with open(p, encoding="utf-8") as fh:
                heads += next(csv.reader(fh))
        elif f.endswith(".xlsx"):
            wb = openpyxl.load_workbook(p, read_only=True)
            for ws in wb.worksheets:
                for r in ws.iter_rows(max_row=12, values_only=True):
                    heads += [str(x) for x in r if x is not None]
    pat = r"forecast|projected|predicted|per (?:broker )?post|formal entr\w* (?:in )?2027|2027 formal|at (?:the )?plan"
    A("R1 no shipped column or cell holds a forecast, a load, or a route at the plan's rates",
      not [h for h in heads if re.search(pat, h, re.I)], [h for h in heads if re.search(pat, h, re.I)][:5])
    T = texts_of(target)
    nums = [(f, m.group(1)) for f, t in T.items() for m in re.finditer(r"([\d,.]+) formal entries a month", t)]
    A("R1 no document states a number of formal entries a month but the standard load (560)",
      nums and all(v == "560" for _, v in nums), nums)
    A("R1 no document states a placement, a load or a forecast of formal entries",
      not [f for f, t in T.items() if re.search(r"posts? (?:placed|go(?:es)?) (?:to|at) (?:Mississauga|Dorval|Richmond|"
                                                r"Calgary|Winnipeg)|(?<!standard )load (?:of|is) \d|forecast (?:of|is) \d",
                                                t)],
      [f for f, t in T.items() if re.search(r"(?<!standard )load (?:of|is) \d|forecast (?:of|is) \d", t)])

    # ------------------------------------------------------------------ R2: every mess item neutral on the call
    est = c["st"]["est"]
    gold_t = tup(c["placed"], c["held"])
    gold_bins = ({g: K.round_to(c["E"][g], 50) for g in GW}, {g: K.round_to(c["final"][g], 5) for g in GW})

    def same(E, label, bins=True):
        p, h, fin, _ = place(E)
        ok = tup(p, h) == gold_t
        if bins:
            ok = ok and ({g: K.round_to(E[g], 50) for g in GW}, {g: K.round_to(fin[g], 5) for g in GW}) == gold_bins
        A(label, ok, (fmt(p, h), {g: round(E[g], 1) for g in GW}))

    both = V.register(target, keep="both")
    seen_wb = defaultdict(list)
    for r in both:
        seen_wb[r["wb"]].append(r)
    twice = {k: v for k, v in seen_wb.items() if len(v) > 1}
    order = list(R.F_HV.values())
    A("R2 amended rows: %d shipments appear twice, the later row in the next quarter's extract and marked amended" %
      len(twice), len(twice) == len(W.amended) and all(
          len(v) == 2 and not v[0]["amended"] and v[1]["amended"] and
          order.index(v[1]["file"]) == order.index(v[0]["file"]) + 1 for v in twice.values()), len(twice))
    n2 = V.with_passthrough([r for r in both if V.in_base(r)], c["rates"], c["line"], lambda k: c["pt"][k][2])
    same(e_of(n2, c["ratio"]), "R2 the amended shipments counted twice: the call and every rounded figure stay (SF11)")
    first = {}
    for r in both:
        if V.in_base(r):
            first.setdefault(r["wb"], r)
    n3 = V.formal_count(list(first.values()), c["rates"], c["line"])
    A("R2 an amended shipment's first row or its amended row: the same route at the plan's rates (SF11)", n3 == c["nb"],
      (n3, c["nb"]))
    def weight_2027(r):
        """A record's 2027 formal entry at the golden's pass-through shares (1 or 0, or the segment's share)."""
        if r["ogd"] or r["cur"] == "CAD":
            return float(r["formal"])
        p_ = pi_w[(r["gw"], r["cur"])]
        return p_ * (K.round_dp(r["val"] * K.PLAN_RATE[r["cur"]], 2) > K.LINE_CAD) + (1 - p_) * r["formal"]
    utc_n = {g: 0.0 for g in GW}
    for r in recs:
        if K.BASE_START <= r["utc"] <= K.BASE_END:
            utc_n[r["gw"]] += weight_2027(r)
    utc_v = {g: sum(1 for s in W.ships if s.gw == g and not (K.BASE_START <= s.utc.date() <= K.BASE_END)) for g in GW}
    A("R2 the UTC release date moves some shipments out of the base window (%d)" % sum(utc_v.values()),
      sum(utc_v.values()) > 0, utc_v)
    same(e_of(utc_n, c["ratio"]), "R2 the base year cut by UTC date: the call and every rounded figure stay (SF12)")
    n4 = defaultdict(int)
    for s in reg:
        if s["ogd"] or V.cents(s["cad"] * c["rates"][s["cur"]] / s["rate"]) > c["line"]:
            n4[s["gw"]] += 1
    A("R2 the CAD value rescaled by the plan's rate over the customs rate: the same counts", dict(n4) == c["nb"],
      (dict(n4), c["nb"]))
    wbp = openpyxl.load_workbook(os.path.join(target, R.F_PLAN27), read_only=True)
    prow = {re.search(r"\((\w{3})\)$", str(r[0])).group(1): r for r in wbp["Inbound volumes"].iter_rows(values_only=True)
            if re.search(r"\((YYZ|YUL|YVR|YYC|YWG)\)$", str(r[0]))}
    ratio_basecol = {g: float(prow[g][13]) / float(prow[g][14]) for g in GW}
    same(e_of(c["n"], ratio_basecol), "R2 the plan's own base column (thousands) as the base: the call and every "
                                      "rounded figure stay")
    ratio_growth = {g: 1 + float(prow[g][15].strip("%+")) / 100.0 for g in GW}
    same(e_of(c["n"], ratio_growth), "R2 the plan's growth column as the scale: the call and every rounded figure stay")
    mm = defaultdict(float)
    plan_k = C.plan_shipments_k(W)
    for r in recs:
        if r["yr"] == 1:
            y, m = r["local"].year, r["local"].month
            mm[r["gw"]] += weight_2027(r) * plan_k[(r["gw"], m)] * 1000.0 / W.volume[(r["gw"], y, m)] / 12.0
    same(dict(mm), "R2 the forecast built month by month on the plan's months: the call and every rounded figure stay "
                   "(SF8)")
    A("R2 the earlier layout's DD/MM/YYYY dates: the count reads no date (the extracts cover the base year), and a "
      "month-first read fails loudly on every REL_DT day above 12",
      any(r["local"].day > 12 for r in reg if r["file"] == R.F_HV["2025Q4"]))
    pgm_none_clvs = sum(1 for s in reg if s["pgm"] == "NONE" and s["route"] == "CLVS")
    A("R2 OGD codes blank or NONE: every row with a program is FORMAL whatever its value, and NONE rows routed CLVS "
      "(%d) rule out reading NONE as a program" % pgm_none_clvs,
      all(s["route"] == "FORMAL" for s in reg if s["ogd"]) and pgm_none_clvs > 1000)
    A("R2 the routes on file are CLR-OPS-11 at each day's customs rate, both years (the extracts' own routing "
      "reproduces)", all((s["route"] == "FORMAL") == (s["ogd"] or s["cad"] > c["line"]) for s in reg_all))
    labels = {g: str(prow[g][0]) for g in GW}
    A("R2 gateways keyed by code, desk name and the plan's region label map one to one",
      all("(%s)" % g in labels[g] for g in GW) and c["st"]["desk"] == K.DESK, labels)
    A("R2 the legacy gateway report (accounting month) is not on the call's path", R.F_GW_LEGACY not in main_reads)
    A("R2 the ask paths' mess sits in files the call never opens", not (main_reads & set(ASK_FILES)))

    # ------------------------------------------------------------------ R3: the analyst test
    Ed = C.default_E(W)
    pd_, hd, fd, od = place(Ed)
    A("R3a the easy route on the clean truth (base-year formal entries per shipment carried on the plan, every filed "
      "rule applied) files one post at each of the five gateways, one held", tup(pd_, hd) == (1, 1, 1, 1, 1, 1),
      fmt(pd_, hd))
    A("R3a the same easy route on the pack files the same", tup(c["placed_d"], c["held_d"]) == (1, 1, 1, 1, 1, 1),
      fmt(c["placed_d"], c["held_d"]))
    Eg = C.golden_E(W, recs)
    pg, hg, fg, og = place(Eg)
    Et = C.truth_E(W)
    ptr, htr, ftr, _ = place(Et)
    A("R3b the world's own 2027 (each shipment routed as its own pricing behaves under the plan's rates, which no file "
      "records) files Dorval 2, Calgary 1, three held", tup(ptr, htr) == (0, 2, 0, 1, 0, 3), fmt(ptr, htr))
    A("R3b the decisive step (each segment's pass-through from the repeat consignments, that share re-routed at the "
      "plan's rates, the rest kept) files the truth's placement and every rounded figure of it",
      tup(pg, hg) == tup(ptr, htr) and {g: K.round_to(Eg[g], 50) for g in GW} == {g: K.round_to(Et[g], 50) for g in GW}
      and {g: K.round_to(fg[g], 5) for g in GW} == {g: K.round_to(ftr[g], 5) for g in GW},
      ({g: round(Eg[g], 1) for g in GW}, {g: round(Et[g], 1) for g in GW}))
    A("R3b the pack's golden equals the truth's", tup(c["placed"], c["held"]) == tup(pg, hg) and
      all(abs(c["E"][g] - Eg[g]) < 1e-6 for g in GW), {g: round(c["E"][g], 3) for g in GW})
    rivals = {"(b) every foreign-currency shipment re-routed at the plan's rates": (C.full_repricing_E(W), "b"),
              "(c) one pass-through share for the network": (C.pooled_share_E(W, recs)[0], "c"),
              "(d) one pass-through share for each invoice currency": (C.by_key_E(W, lambda r: r["cur"], recs)[0],
                                                                       "dc")}
    for name, (E_, k_) in rivals.items():
        p_, h_, f_, _ = place(E_)
        A("R3c rival %s files Dorval 2 and four held, Calgary's post lost at a load of %.1f (truth and pack)" %
          (name, E_["YYC"] / 4), tup(p_, h_) == (0, 2, 0, 0, 0, 4) == tup(c["placed_" + k_], c["held_" + k_]) and
          E_["YYC"] / 4 <= 550 and all(abs(c["E" + k_][g] - E_[g]) < 1e-6 for g in GW), (fmt(p_, h_), round(E_["YYC"], 1)))
    pt = c["pt"]
    a_, b_, p_yyc = pt[("YYC", "USD")]
    others = {k: v for k, v in pt.items() if k != ("YYC", "USD")}
    cad_held = sorted(k for k, v in pt.items() if v[2] <= 0.05)
    A("R3c the segments that keep their CAD values are Calgary's US-dollar goods and the Chinese goods at Mississauga and "
      "Richmond, each 0.02 or under; every other segment 0.95 or over", cad_held == [("YVR", "CNY"), ("YYC", "USD"),
                                                                                ("YYZ", "CNY")] and
      all(v[2] <= 0.02 for k, v in pt.items() if k in cad_held) and
      all(v[2] >= 0.95 for k, v in pt.items() if k not in cad_held), {k: round(v[2], 3) for k, v in pt.items()})
    A("R3c the data rejects (b): Calgary's US-dollar goods repeat a consignment of the year before at the same CAD value "
      "%d times and at the same invoice value %d times (share re-routed %.3f)" % (b_, a_, p_yyc),
      p_yyc <= 0.02 and b_ >= 10000, (a_, b_, p_yyc))
    usd4 = [pt[(g, "USD")][2] for g in ("YYZ", "YUL", "YVR", "YWG")]
    Egw, pi_gw = C.by_key_E(W, lambda r: r["gw"], recs)
    pgw, hgw, fgw, _ = place(Egw)
    A("R3c SF21 rival (e) one pass-through share for each gateway files the golden's placement on other figures (Mississauga "
      "%d, Richmond %d a month), rejected by the data: Mississauga's and Richmond's Chinese goods keep their CAD values "
      "(%.3f, %.3f) and their US-dollar goods their invoice values (%.3f, %.3f)" % (
          K.round_to(Egw["YYZ"], 50), K.round_to(Egw["YVR"], 50), pt[("YYZ", "CNY")][2], pt[("YVR", "CNY")][2],
          pt[("YYZ", "USD")][2], pt[("YVR", "USD")][2]),
      tup(pgw, hgw) == gold_t and ({g: K.round_to(Egw[g], 50) for g in GW}, {g: K.round_to(fgw[g], 5) for g in GW})
      != gold_bins and max(pt[("YYZ", "CNY")][2], pt[("YVR", "CNY")][2]) <= 0.02 and
      min(pt[("YYZ", "USD")][2], pt[("YVR", "USD")][2]) >= 0.99, {g: round(Egw[g], 1) for g in GW})
    inv_held = {k: v for k, v in others.items() if k[1] != "CNY"}
    A("R3c SF21 the data rejects (c) and (d): the US-dollar, euro, sterling and yen goods keep their invoice values (shares "
      "%.3f to %.3f; the other four gateways' US-dollar goods %.3f or more) against Calgary's %.3f, where one network "
      "share is %.3f and one US-dollar share %.3f" % (min(v[2] for v in inv_held.values()),
                                                      max(v[2] for v in inv_held.values()), min(usd4), p_yyc,
                                                      c["ptc"]["all"][2], c["ptd"]["USD"][2]),
      min(v[2] for v in inv_held.values()) >= 0.95 and min(usd4) >= 0.97 and 0.6 <= c["ptc"]["all"][2] <= 0.9 and
      0.6 <= c["ptd"]["USD"][2] <= 0.9)
    rows_pt = list(pt.values())
    p_all = sum(a for a, b, q in rows_pt) / sum(a + b for a, b, q in rows_pt)
    chi = sum((a - (a + b) * p_all) ** 2 / ((a + b) * p_all) + (b - (a + b) * (1 - p_all)) ** 2 / ((a + b) * (1 - p_all))
              for a, b, q in rows_pt)
    se_yyc = math.sqrt(p_yyc * (1 - p_yyc) / (a_ + b_))
    A("R3c the margins: homogeneity across the %d segments %.0f on %d degrees of freedom; Calgary's share %.4f plus or "
      "minus %.4f (95 per cent), %.0f standard errors from one and %.0f under the network's %.3f" % (
          len(rows_pt), chi, len(rows_pt) - 1, p_yyc, 1.96 * se_yyc, (1 - p_yyc) / se_yyc, (p_all - p_yyc) / se_yyc,
          p_all), chi > 10000 and (1 - p_yyc) / se_yyc > 500 and (p_all - p_yyc) / se_yyc > 500 and
      abs(p_all - c["ptc"]["all"][2]) < 1e-12, round(chi))
    rungs = {
        "R0b trend": ({g: C.trend_E(W, g) for g in GW}, (1, 1, 1, 1, 1, 1)),
        "R0c smoothing": ({g: C.holt_winters_E(W, g) for g in GW}, (2, 1, 1, 1, 1, 0)),
        "R0d last quarter": ({g: C.latest_quarter_E(W, g) for g in GW}, (1, 1, 1, 1, 1, 1)),
        "R0e pooled formal share": (C.pooled_E(W), (3, 1, 1, 0, 1, 0)),
        "R0f last customs rate": (C.spot_E(W, recs), (1, 1, 1, 1, 1, 1)),
        "R1 share fitted on the rates": ({g: C.share_regression_E(W, g) for g in GW}, (1, 1, 1, 1, 0, 2)),
        "R2 every shipment re-routed": (C.full_repricing_E(W), (0, 2, 0, 0, 0, 4)),
        "R2a one network share": (C.pooled_share_E(W, recs)[0], (0, 2, 0, 0, 0, 4)),
        "R2b one share per currency": (C.by_key_E(W, lambda r: r["cur"], recs)[0], (0, 2, 0, 0, 0, 4)),
        "R2c US dollar's change alone": (C.usd_only_E(W, recs), (0, 1, 0, 1, 0, 4)),
        "R2d euro shipments alone": (C.eur_only_E(W, recs), (1, 2, 1, 1, 1, 0)),
        "R2e value alone, OGD ignored": (C.value_only_E(W, recs), (0, 0, 0, 1, 0, 5)),
    }
    res = {}
    for name, (E, want) in rungs.items():
        p, h, _, _ = place(E)
        res[name] = (E, tup(p, h))
        A("R3b rung %s files %s" % (name, want), tup(p, h) == want, fmt(p, h))
        A("R3b rung %s stays wrong (two or more buckets off the golden)" % name,
          sum(1 for a, b in zip(tup(p, h), tup(pg, hg)) if a != b) >= 2, tup(p, h))
    A("R3b the default differs from the golden in five of six buckets",
      sum(1 for a, b in zip(tup(pd_, hd), tup(pg, hg)) if a != b) == 5)
    A("R3b every construction that keeps the base year's routing staffs Richmond",
      all(res[n][1][2] >= 1 for n in ("R0b trend", "R0c smoothing", "R0d last quarter", "R0e pooled formal share",
                                      "R0f last customs rate")) and pd_["YVR"] == 1)
    A("R3b every construction that re-routes Calgary's goods at the plan's US dollar rate loses Calgary's post",
      all(res[n][1][3] == 0 for n in ("R2 every shipment re-routed", "R2a one network share",
                                      "R2b one share per currency")))
    blocks = Counter()
    for f in R.F_HV.values():
        with open(os.path.join(target, f), encoding="utf-8") as fh:
            rr = [tuple(r[i] for i in (1, 5, 7, 8)) for r in csv.reader(fh)][1:]
        for i in range(len(rr) - 4):
            blocks[tuple(rr[i:i + 5])] += 1
    A("R3c no copied block: no five-row run of the register repeats", max(blocks.values()) == 1, max(blocks.values()))
    rep = defaultdict(list)
    with open(os.path.join(target, R.F_GW_NEW), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            rep[r["gateway_code"]].append((int(r["inbound_international_shipments"]), int(r["formal_entries_released"])))
    A("R3c no gateway's month repeats another's in the gateway report",
      all(len(set(v)) == len(v) for v in rep.values()) and
      all(len({rep[g][i][1] for g in GW}) == 5 for i in range(len(rep["YYZ"]))))
    A("R3c no memo, draft, briefing or stakeholder figure: the notes and documents state no forecast, load or placement",
      not [f for f, t in T.items() if re.search(r"recommend|we propose|should (?:go|be placed)|forecast for 2027", t,
                                                re.I)])

    # ------------------------------------------------------------------ the ladder's arithmetic
    steps = [(g, round(l, 1)) for g, l in og]
    A("ladder: the placement steps are Dorval, Calgary, Dorval", [g for g, _ in og] == ["YUL", "YYC", "YUL"], steps)
    A("ladder: every placed post clears 560 by 7 per cent or more", min(l for _, l in og) / 560 >= 1.07, steps)
    top = max(fg, key=fg.get)
    A("ladder: the stop is Mississauga, 4 per cent or more under 560", top == "YYZ" and fg[top] <= 0.96 * 560,
      (top, round(fg[top], 1)))
    A("ladder: flip conditions, Mississauga needs 3,920 a month for a fourth post and Dorval keeps its second above "
      "3,360, Calgary its post above 2,240", Eg["YYZ"] < 560 * 7 and Eg["YUL"] > 560 * 6 and Eg["YYC"] > 560 * 4 and
      abs(560 * 7 - 3920) < 1e-9, (Eg["YYZ"], Eg["YUL"], Eg["YYC"]))
    l0 = {g: Ed[g] / K.ESTABLISHMENT[g] for g in GW}
    rank = sorted(GW, key=lambda g: -l0[g])
    A("ladder: Dorval, the golden's two-post desk, is second of five on the default's loads before any post",
      rank.index("YUL") == 1, [(g, round(l0[g], 1)) for g in rank])
    mv = {g: Eg[g] / Ed[g] - 1 for g in GW}
    A("ladder: against the default the decisive step raises Dorval by 10 per cent or more, cuts Mississauga, Richmond "
      "and Winnipeg by 10 per cent or more and keeps Calgary within 2 per cent",
      mv["YUL"] >= 0.10 and all(mv[g] <= -0.10 for g in ("YYZ", "YVR", "YWG")) and abs(mv["YYC"]) <= 0.02,
      {g: round(100 * v, 1) for g, v in mv.items()})
    Eb_ = res["R2 every shipment re-routed"][0]
    mb = {g: Eg[g] / Eb_[g] - 1 for g in GW}
    A("ladder: against every shipment re-routed the pass-through raises Calgary by 25 per cent or more and Mississauga "
      "and Richmond by 0.5 per cent or more (their Chinese goods keep their CAD values), and moves Dorval and Winnipeg "
      "by under 0.1 per cent", mb["YYC"] >= 0.25 and mb["YYZ"] >= 0.005 and mb["YVR"] >= 0.005 and
      all(abs(mb[g]) < 0.001 for g in ("YUL", "YWG")), {g: round(100 * v, 2) for g, v in mb.items()})
    eb_b = {g: (K.round_to(Eb_[g], 50), K.round_to(Eb_[g] / (K.ESTABLISHMENT[g] + place(Eb_)[0][g]), 5)) for g in GW}
    eb_g = {g: (K.round_to(Eg[g], 50), K.round_to(fg[g], 5)) for g in GW}
    A("ladder: every shipment re-routed misses Calgary's, Mississauga's and Richmond's forecast and load and the held "
      "count, and keeps Dorval's and Winnipeg's", [g for g in GW if eb_b[g] != eb_g[g]] == ["YYZ", "YVR", "YYC"],
      (eb_b, eb_g))
    low = {n: res[n][0]["YUL"] / Ed["YUL"] - 1 for n in ("R0b trend", "R0c smoothing", "R0d last quarter",
                                                          "R0f last customs rate", "R1 share fitted on the rates")}
    A("ladder: dominance, no rung that keeps the history's rates moves Dorval by half the decisive step's move",
      max(abs(v) for v in low.values()) <= 0.5 * mv["YUL"], {n: round(100 * v, 1) for n, v in low.items()})

    # ------------------------------------------------------------------ the semantic fork table
    hr = openpyxl.load_workbook(os.path.join(target, R.F_HR_STAFF), read_only=True).worksheets[0]
    hrr = list(hr.iter_rows(values_only=True))
    ix = {h: i for i, h in enumerate(hrr[0])}
    heads_desk = Counter(r[ix["Desk"]] for r in hrr[1:] if r[ix["Status"]] == "Active")
    A("SF1 the establishment at 1 January 2027 equals the brokers each desk holds at the pull",
      {g: heads_desk[K.DESK[g]] for g in GW} == est, dict(heads_desk))
    hvh = []
    for f in R.F_HV.values():
        with open(os.path.join(target, f), encoding="utf-8") as fh:
            hvh += next(csv.reader(fh))
    A("SF2 the register carries no transmitter, so 'whoever transmits it' leaves one count",
      not [h for h in hvh if re.search(r"broker|transmit|filer|agent", h, re.I)])
    A("SF3 BRK-STD-03 4.2 files the average month of the plan year", "the average month of its forecast formal entries "
      "in the plan year" in T[R.F_STD])
    p4 = {g: 0 for g in GW}
    for _ in range(K.NEW_POSTS):
        ld = {g: Eg[g] / (est[g] + p4[g] + 1) for g in GW}
        g = max(ld, key=ld.get)
        if ld[g] > 560:
            p4[g] += 1
        else:
            break
    A("SF4 counting the post being placed files Dorval 1, five held, ruled out by 4.5", tup(p4, 6 - sum(p4.values())) ==
      (0, 1, 0, 0, 0, 5) and "counting the posts already placed in this review" in T[R.F_STD], fmt(p4, 6 - sum(p4.values())))
    gaps = [abs(l - 560) for _, l in og] + [abs(max(fg.values()) - 560)]
    A("SF5 no load at a placement step or the stop within 25 of 560 (exceeds and at-or-above agree)", min(gaps) >= 25,
      [round(x, 1) for x in gaps])
    A("SF6 the last customs rate on file (the golden's pass-through at it) files a post at every gateway (BRK-STD-03 4.3 "
      "and the plan's assumption 7 rule it out)", res["R0f last customs rate"][1] == (1, 1, 1, 1, 1, 1) and "the plan's planning assumptions" in
      T[R.F_STD] and "Treasury planning rates for 2027" in T[R.F_PLAN27])
    A("SF7 the plan files its base, the twelve months October 2025 to September 2026, and the profile as in the base",
      "The twelve months October 2025 to September 2026." in T[R.F_PLAN27] and
      "Each gateway's lanes, customers and goods as in the base." in T[R.F_PLAN27])
    two = {}
    pk = C.plan_shipments_k(W)
    for g in GW:
        n2y = sum(weight_2027(r) for r in recs if r["gw"] == g)
        v2y = sum(W.volume[(g, y, m)] for (y, m) in K.PRIOR_MONTHS + K.BASE_MONTHS)
        two[g] = n2y / v2y * sum(pk[(g, m)] for m in range(1, 13)) * 1000 / 12
    p2y, h2y, _, _ = place(two)
    A("SF7 the two years' profile instead of the base's (ruled out by assumption 1 and 2): the same call and the same "
      "monthly forecasts to the nearest 50", tup(p2y, h2y) == gold_t and
      {g: K.round_to(two[g], 50) for g in GW} == gold_bins[0], {g: round(two[g], 1) for g in GW})
    ok8 = all(abs(plan_k[(g, m)] - W.volume[(g, y, m)] * (1 + K.PLAN_GROWTH[g]) / 1000.0) <= 0.05 + 1e-9
              for g in GW for (y, m) in K.BASE_MONTHS)
    A("SF8 the plan's months are the base months grown by the plan's growth, to 0.1 thousand", ok8)
    for f_ in (0.98, 0.99, 1.01, 1.02, 1.03):
        p, h, _, _ = place(C.inflated_E(W, f_, recs))
        A("SF9 every price moved by %+.0f per cent before 2027 (invoice-currency and CAD prices alike): the placement "
          "stays" % (100 * (f_ - 1)), tup(p, h) == gold_t, fmt(p, h))
    A("SF10 OGD goods routed on value alone hold five posts, ruled out by CLR-OPS-11 3.1 (b)",
      res["R2e value alone, OGD ignored"][1] == (0, 0, 0, 1, 0, 5) and "whatever its value" in T[R.F_CLR])
    A("SF11 asserted above (amended rows)", True)
    A("SF12 asserted above (UTC dates)", True)
    worst = max(K.PLAN_RATE[s.cur] / s.rate for s in W.ships if s.cur != "CAD")
    floor_ok = all(s.ogd or s.vfd_cad >= K.CHECK_FLOOR_CAD for s in W.ships)
    A("SF13 nothing under CAD 2,500 can cross CAD 3,300 at a plan rate: the largest rise is %.1f per cent" %
      (100 * (worst - 1)), worst < K.LINE_CAD / K.CHECK_FLOOR_CAD and floor_ok, worst)
    near = min(abs(V.cents(s["val"] * c["rates"][s["cur"]]) - c["line"]) for s in reg if not s["ogd"])
    A("SF14 no value lands within CAD 0.75 of the line at the plan's rates (over and at-or-over agree)", near >= 0.75,
      near)
    A("SF14 no CAD value on file sits at CAD 3,300.00, so the routes kept for the CAD-held share read alike over and "
      "at-or-over", not [s for s in reg_all if not s["ogd"] and s["cad"] == c["line"]])
    A("SF15 every shipment invoiced in CAD carries a rate of 1", all(s["rate"] == 1.0 for s in reg if s["cur"] == "CAD"))
    for step in (1, 5):
        Er = {g: K.round_to(Eg[g], step) for g in GW}
        p, h, _, _ = place(Er)
        A("SF16 the forecast rounded to %d before placing: the same call" % step, tup(p, h) == gold_t, fmt(p, h))
    close = []
    pl = {g: 0 for g in GW}
    for _ in range(K.NEW_POSTS):
        ld = sorted((Eg[g] / (est[g] + pl[g]) for g in GW), reverse=True)
        close.append(ld[0] - ld[1])
        g = max(GW, key=lambda x: Eg[x] / (est[x] + pl[x]))
        if Eg[g] / (est[g] + pl[g]) <= 560:
            break
        pl[g] += 1
    A("SF17 no two loads within 5 of each other at any placement step", min(close) >= 5, [round(x, 1) for x in close])
    A("SF18 the plan's volumes are thousands (its heading), and thousands times 1,000 match the report's base",
      "Inbound international shipments, thousands" in T[R.F_PLAN27] and
      all(abs(float(prow[g][14]) * 1000 - c["rep"][g][0]) <= 50 for g in GW))

    # ------------------------------------------------------------------ the decisive inference: each segment's pass-through
    slopes = C.slope_shares(recs)
    held3 = (("YYC", "USD"), ("YYZ", "CNY"), ("YVR", "CNY"))
    sl_o = [v for k, v in slopes.items() if k not in held3]
    A("R3 pass-through: within the consignments seen in both years the CAD value moves with the day's rate in every "
      "segment but three (slopes %.3f to %.3f) and holds in Calgary's US-dollar goods and the Chinese goods at "
      "Mississauga and Richmond (slopes %s)" % (min(sl_o), max(sl_o), ", ".join("%.3f" % slopes[k] for k in held3)),
      min(sl_o) >= 0.95 and all(slopes[k] <= 0.02 for k in held3), slopes)
    si_cp = [i for i, sg in enumerate(K.SEGMENTS["YYC"]) if sg["kind"] == "pp"][0]
    cps = sorted(v for v, _ in K.SEGMENTS["YYC"][si_cp]["cp"])
    yyc = [s for s in reg_all if s["gw"] == "YYC" and s["cur"] == "USD" and not s["ogd"]]
    ok_cp, moves = True, []
    r_prior = [s["rate"] for s in yyc if K.PRIOR_START <= s["local"] <= K.PRIOR_END]
    r_base = [s["rate"] for s in yyc if V.in_base(s)]
    rate_move = (sum(r_base) / len(r_base)) / (sum(r_prior) / len(r_prior)) - 1
    for P in cps:
        at = [s for s in yyc if K.round_to(s["cad"], 1) == P and abs(s["cad"] - P) <= 0.011]
        pv = [s["val"] for s in at if K.PRIOR_START <= s["local"] <= K.PRIOR_END]
        bv = [s["val"] for s in at if V.in_base(s)]
        ok_cp = ok_cp and len(pv) >= 20 and len(bv) >= 20 and len(set(bv)) >= 10
        if pv and bv:
            moves.append((sum(bv) / len(bv)) / (sum(pv) / len(pv)) - 1)
    A("R3 pass-through: each of Calgary's %d CAD price points (CAD %d to %d) recurs in both years at the same CAD value "
      "while its US-dollar invoice values move against the rate (invoice values %+.2f to %+.2f per cent, the rate "
      "%+.2f)" % (len(cps), cps[0], cps[-1], 100 * min(moves), 100 * max(moves), 100 * rate_move),
      ok_cp and all(abs(m + rate_move) <= 0.004 for m in moves), [round(100 * m, 2) for m in moves])
    lists_ok, drift = True, {}
    for g in GW:
        for si, seg in enumerate(K.SEGMENTS[g]):
            if seg["kind"] != "pp" or K.lcp_share(g, seg["cur"]) >= 0.5:
                continue
            lp = {v for v, _ in seg["p"]}
            pv = {s["val"] for s in reg_all if s["gw"] == g and s["cur"] == seg["cur"] and s["val"] in lp and
                  K.PRIOR_START <= s["local"] <= K.PRIOR_END}
            bv = {s["val"] for s in reg if s["gw"] == g and s["cur"] == seg["cur"] and s["val"] in lp}
            lists_ok = lists_ok and pv == bv == lp
            pc = [s["cad"] / s["val"] for s in reg_all if s["gw"] == g and s["cur"] == seg["cur"] and s["val"] in lp
                  and K.PRIOR_START <= s["local"] <= K.PRIOR_END]
            bc = [s["cad"] / s["val"] for s in reg if s["gw"] == g and s["cur"] == seg["cur"] and s["val"] in lp]
            drift[(g, seg["cur"])] = round(100 * (sum(bc) / len(bc) / (sum(pc) / len(pc)) - 1), 1)
    A("R3 pass-through: every EUR, USD and GBP list price outside Calgary recurs in both years at its invoice value while "
      "its CAD value moves with the rate (%s)" % ", ".join("%s %s %+.1f%%" % (g, cu, d) for (g, cu), d in drift.items()),
      lists_ok and {cu for _, cu in drift} >= {"EUR", "USD", "GBP"} and max(abs(d) for d in drift.values()) >= 3.0,
      drift)
    A("SF19 no document files how the goods are priced, how an exchange rate passes into a value, or which goods hold "
      "their CAD values (the plan's profile names lanes, customers and goods only)",
      not [f for f, t in T.items() if re.search(PRICE_WORDS, t, re.I)], [f for f, t in T.items() if re.search(PRICE_WORDS, t, re.I)])

    # ------------------------------------------------------------------ SF20: every defensible pass-through estimator
    nohs_i = lambda r: (r["gw"], r["imp"], r["cur"], r["val"])
    nohs_c = lambda r: (r["gw"], r["imp"], r["cur"], r["cad"])
    dol_c = lambda r: (r["gw"], r["imp"], r["hs"], r["cur"], K.round_to(r["cad"], 1))

    def by_counts(**kw):
        return C.scale(W, C.mixture_counts(recs, C.shares(*C.repeat_counts(recs, **kw))))
    dol_nohs = lambda r: (r["gw"], r["imp"], r["cur"], K.round_to(r["cad"], 1))
    p0 = [r for r in recs if r["yr"] == 0 and not r["ogd"] and r["cur"] != "CAD"]
    p0_inv, p0_dol = {C.k_inv(r) for r in p0}, {dol_nohs(r) for r in p0}
    chance = sum(1 for r in recs if r["yr"] == 1 and r["gw"] == "YYZ" and not r["ogd"] and r["cur"] != "CAD" and
                 not r["cad_priced"] and C.k_inv(r) not in p0_inv and dol_nohs(r) in p0_dol)
    A("SF20 the crude keys left out: a CAD value to the whole dollar without the tariff item matches %d of Mississauga's "
      "base-year shipments priced in their invoice currency to a shipment of the year before by chance" % chance,
      chance >= 1000, chance)
    fam = {
        "shares counted by consignment": by_counts(weight="consignment"),
        "both years matched each way": by_counts(both_years=True),
        "the consignment key without the tariff item": by_counts(ki=nohs_i, kc=nohs_c),
        "shares counted by consignment, the key without the tariff item": by_counts(weight="consignment", ki=nohs_i,
                                                                                   kc=nohs_c),
        "the CAD value matched within a cent": by_counts(tol=True),
        "shares counted by consignment, the CAD value within a cent": by_counts(tol=True, weight="consignment"),
        "the CAD value matched to the dollar": by_counts(kc=dol_c),
        "the within-consignment regression of the log CAD value on the log rate":
            C.scale(W, C.mixture_counts(recs, C.slope_shares(recs))),
        "the regression of the change in the log CAD value on the change in the log rate":
            C.scale(W, C.mixture_counts(recs, C.diff_slope_shares(recs))),
        "each repeat routed as its own consignment behaves": C.consignment_level_E(W, recs),
        "the world's truth": Et,
    }
    for name, E_ in fam.items():
        p_, h_, f_, _ = place(E_)
        A("SF20 the pass-through estimated with %s: the same call and every rounded figure" % name, tup(p_, h_) == gold_t and
          ({g: K.round_to(E_[g], 50) for g in GW}, {g: K.round_to(f_[g], 5) for g in GW}) == gold_bins,
          ({g: round(E_[g], 1) for g in GW}, {g: round(f_[g], 2) for g in GW}))

    # ------------------------------------------------------------------ rounding: per post from either forecast
    post_after = {g: K.ESTABLISHMENT[g] + pg[g] for g in GW}
    for g in GW:
        x = Eg[g] / post_after[g]
        y = K.round_to(Eg[g], 50) / post_after[g]
        A("rounding: %s entries per post %d from the unrounded forecast (%.2f, 1.0 or more from its edge) and from the "
          "forecast to the nearest 50 (%.2f, fixed by that bin, 0.5 or more)" % (K.DESK[g], K.round_to(x, 5), x, y),
          K.round_to(x, 5) == K.round_to(y, 5) and edge(x, 5) >= 1.0 and edge(y, 5) >= 0.5,
          (round(edge(x, 5), 2), round(edge(y, 5), 2)))
    st_loads = [(g, l) for g, l in og]
    A("rounding: each placement step's load sits 1.0 or more from its edge (%s)" %
      ", ".join("%s %.1f" % (K.DESK[g], l) for g, l in st_loads), all(edge(l, 5) >= 1.0 for _, l in st_loads))

    # ------------------------------------------------------------------ bins
    for g in GW:
        A("bin: %s forecast %.1f sits %.1f from its 50 edge (15 or more)" % (K.DESK[g], Eg[g], edge(Eg[g], 50)),
          edge(Eg[g], 50) >= 15.0)
        A("bin: %s entries per post %.2f sits %.2f from its 5 edge (1.5 or more)" % (K.DESK[g], fg[g], edge(fg[g], 5)),
          edge(fg[g], 5) >= 1.5)
    Ag, Bg = G["A"], G["B"]
    A("bin: every overtime figure sits 0.15 hour or more from its edge",
      min(edge(v, 1) for v in Ag.values()) >= 0.15, round(min(edge(v, 1) for v in Ag.values()), 3))
    A("bin: no Brandmoor figure on a half dollar", min(edge(v, 1) for v in Bg.values()) >= 0.04,
      round(min(edge(v, 1) for v in Bg.values()), 3))

    # ------------------------------------------------------------------ the asks
    A("ask A: the pack's figures equal the truth's to the hour",
      all(K.round_to(Ap[(K.DESK[g], q)], 1) == K.round_to(Ag[(g, q)], 1) and abs(Ap[(K.DESK[g], q)] - Ag[(g, q)]) < 0.01
          for g in GW for q in (1, 2, 3, 4)))
    A("ask B: the pack's figures equal the truth's to the cent",
      all(abs(Bp[(K.DESK[g], q)] - Bg[(g, q)]) < 0.005 for g in GW for q in (1, 2, 3, 4)))
    lines_v = V.time_lines(target, V.staff(target))
    mig = sorted((l["emp"], l["day"], l["code"], l["hours"], round(l["worked"], 4)) for l in lines_v
                 if l["src"] == "MIGRATED")
    leg = sorted((l["emp"], l["day"], l["code"], l["hours"], round(l["worked"], 4)) for l in lines_v
                 if l["src"] == "LEGACY" and FD.TA_NEW_FROM <= l["day"] <= FD.TA_LEGACY_END)
    A("ask A: the migrated week stands in both exports, line for line, under two key schemes (%d lines)" % len(mig),
      mig == leg and len(mig) > 20, (len(mig), len(leg)))
    ot_punch = all(abs(l["hours"] - l["worked"]) < 1e-6 for l in lines_v if l["code"] == "OT")
    A("ask A: every overtime line's hours equal its punches (the hours and the punches agree)", ot_punch)
    cb_short = sum(1 for l in lines_v if l["code"] == "CB" and l["worked"] < 3 and l["hours"] == 3.0)
    A("ask A: call-back lines pay the three-hour minimum where shorter (%d lines)" % cb_short, cb_short > 30)
    for mode, label, want in (("s1", "S1 the home cost centre the export stamped", "all"),
                              ("s2", "S2 call-backs at their paid hours", "all"),
                              ("s3", "S3 the migration week twice", "fq3"),
                              ("cracker", "S4 the cracker: all three", "all"),
                              ("s5", "S5 the HR file's current desk", "all")):
        S = C.ask_a(F.staff, F.lines, mode)
        diffs = {k: S[k] - Ag[k] for k in Ag}
        if want == "all":
            ok = all(abs(v) >= 1.0 and K.round_to(S[k], 1) != K.round_to(Ag[k], 1) for k, v in diffs.items())
        else:
            ok = all((abs(v) >= 1.0) if k[1] == 3 else (abs(v) < 1e-9) for k, v in diffs.items())
        A("ask A stop %s misses %s" % (label, "every figure" if want == "all" else "every October to December figure"),
          ok, round(min(abs(v) for k, v in diffs.items() if want == "all" or k[1] == 3), 2))
    movers = [s for s in F.staff if len(s.desks) > 1]
    A("ask A: seven brokers moved desks and every desk sends or receives a mover", len(movers) == 7 and
      {g for s in movers for _, g in s.desks} == set(GW))
    for mode, label, want in (("s1", "S1 the AP register by posting date (the cracker)", "all"),
                              ("s2", "S2 the statements by their titled month", "all"),
                              ("s3", "S3 the portal by transmission, the adjustment on every line", "all"),
                              ("s4", "S4 the portal by transmission, the adjustment once", "all"),
                              ("s5", "S5 the disputed original kept beside its re-issue", "yyz4")):
        S = C.ask_b(F.stmts, mode)
        diffs = {k: S[k] - Bg[k] for k in Bg}
        if want == "all":
            ok = all(abs(v) >= 1.0 for v in diffs.values())
        else:
            ok = all((abs(v) >= 1.0) if k == ("YYZ", 4) else abs(v) < 1e-9 for k, v in diffs.items())
        A("ask B stop %s misses %s" % (label, "every figure" if want == "all" else "Mississauga's January to March"),
          ok, round(min(abs(v) for k, v in diffs.items() if want == "all" or k == ("YYZ", 4)), 2))
    pm = V.portal_months(target)
    A("ask B: every statement in the portal bills the month before its statement month",
      all(len(v) == 1 for v in pm.values()))
    ap = V.ap_register(target)
    A("ask B: the AP register posts each statement in the month it is dated (the register runs a month behind the "
      "transmissions)", all((r["post"].year, r["post"].month) == (r["date"].year, r["date"].month) for r in ap
                            if r["type"] == "Invoice"))
    A("ask B: the credit memo reversing the disputed statement is dated after the year end",
      [r["date"] for r in ap if r["type"] == "Credit memo"] == [dt.date(2026, 4, 6)])
    prt = V.portal_lines(target)
    A("ask B: the portal holds no March 2026 transmission", not [r for r in prt if r["transmitted_et"][:7] == "2026-03"])
    A("ask B: the portal repeats each statement's minimum adjustment on every line",
      all(len({r["statement_minimum_adjustment"] for r in prt if r["statement_number"] == n}) == 1 for n in pm) and
      any(float(r["statement_minimum_adjustment"]) > 0 for r in prt))
    A("ask B: the AP register runs to September 2026 postings", max(r["post"] for r in ap) >= dt.date(2026, 9, 1))
    re_ = [s_ for s_ in F.stmts if s_.reissue_of]
    org = [s_ for s_ in F.stmts if s_.credit_memo]
    A("ask B: the golden's note matches the dispute (Mississauga, February 2026, re-issued 23 March, memo 6 April)",
      len(re_) == 1 and len(org) == 1 and re_[0].gw == "YYZ" and re_[0].month == (2026, 2) and
      re_[0].issued == dt.date(2026, 3, 23) and org[0].credit_memo[1] == dt.date(2026, 4, 6))
    A("ask B: Brandmoor's fiscal 2025-26 total is not round", sum(Bg.values()) % 100 > 1, round(sum(Bg.values()), 2))
    title = DL.chart_title(G)
    peaks = {g: max((1, 2, 3, 4), key=lambda q: Ag[(g, q)]) for g in GW}
    tops = {q: max(GW, key=lambda g: Ag[(g, q)]) for q in (1, 2, 3, 4)}
    A("ask A: the overtime slide's title is true (%s)" % title,
      (title.startswith("Overtime peaked at every desk") and all(p == 3 for p in peaks.values())) or
      (title.startswith(K.DESK[tops[1]]) and len(set(tops.values())) == 1))
    A("asks: the separation line, the asks read no file the call reads",
      not ((a_reads | b_reads) & set(MAIN_FILES)), sorted((a_reads | b_reads) & set(MAIN_FILES)))
    A("asks: no filed sentence says when a home cost centre was stamped or which month a statement bills",
      not re.search(r"stamped|as of the (?:pull|export)|previous month|month before|preceding month|bills the",
                    T[R.F_NOTE_PAY] + T[R.F_BCS_RATES] + T[R.F_BCS_PDF], re.I))

    # ------------------------------------------------------------------ deletion tests
    d1 = copy_without(target, set(dis))
    d2 = copy_without(target, set(ASK_FILES))
    d3 = copy_without(target, set(MAIN_FILES))
    try:
        c1 = V.main_call(d1)
        A("distractors: deleting both leaves the call and every figure", tup(c1["placed"], c1["held"]) == gold_t and
          c1["E"] == c["E"] and V.ask_a(d1) == Ap and V.ask_b(d1) == Bp)
        c2 = V.main_call(d2)
        A("separation: the call computed with every ask file removed is the golden", c2["E"] == c["E"] and
          tup(c2["placed"], c2["held"]) == gold_t)
        A("separation: the asks computed with every main-call file removed are the asks",
          V.ask_a(d3) == Ap and V.ask_b(d3) == Bp)
    finally:
        for d in (d1, d2, d3):
            shutil.rmtree(d)
    t26 = T[R.F_PLAN26]
    A("distractors: the 2026 plan is superseded by the 2027 plan (other year, other rates) and states no formal entries",
      "Network plan 2026" in t26 and "1.6800" not in t26 and "formal" not in t26.lower())
    with open(os.path.join(target, R.F_EXPORTS), encoding="utf-8") as fh:
        eh = next(csv.reader(fh))
    A("distractors: the export declaration report counts outbound declarations, not import entries",
      "export_declarations_filed" in eh and not [h for h in eh if "formal" in h])

    # ------------------------------------------------------------------ load-bearing facts once, no connecting sentence
    def where(pat):
        return sorted(f for f, t in T.items() if re.search(pat, t))
    p3300 = r"(?<![\d,.])3,?300(?![.,]?\d)"
    p560 = r"(?<![\d,.])560(?![.,]?\d)"
    A("facts: the CAD 3,300 line is filed in CLR-OPS-11 alone", where(p3300) == [R.F_CLR], where(p3300))
    A("facts: the standard load 560 is filed in BRK-STD-03 alone", where(p560) == [R.F_STD], where(p560))
    A("facts: the 2027 planning rates are filed in the 2027 plan alone", where(r"1\.6800") == [R.F_PLAN27],
      where(r"1\.6800"))
    A("facts: the establishment is filed in BRK-STD-03 appendix A alone",
      where(r"Establishment at 1 January 2027") == [R.F_STD])
    A("facts: no document connects the plan's rates to routing or to formal entries (BRK-STD-03 names no exchange rate, "
      "CLR-OPS-11 no plan, the plan no entry or route)",
      not re.search(r"exchange|currenc", T[R.F_STD], re.I) and not re.search(r"\bplan\b|planning", T[R.F_CLR], re.I) and
      not re.search(r"formal|entr(?:y|ies)|rout|brokerage", T[R.F_PLAN27], re.I))

    # ------------------------------------------------------------------ realism sweep
    A("realism: the register's shipment count is not round", len(reg) % 100 != 0, len(reg))
    A("realism: no gateway's base-year shipments are round", all(c["rep"][g][0] % 100 for g in GW))
    A("realism: the plan's network total is not round", abs(sum(sum(c["vols"][g]) for g in GW)) % 10 > 0.05)
    daily = defaultdict(Counter)
    for s in reg:
        daily[s["gw"]][s["local"]] += 1
    A("realism: no constant stream, every gateway's daily check count varies",
      all(len(set(v.values())) > 20 for v in daily.values()))
    banned = (b"openpyxl", b"python-pptx", b"ReportLab", b"reportlab", b"matplotlib", b"Matplotlib")
    leaky = []
    for folder in (target, golden):
        for f in os.listdir(folder):
            p = os.path.join(folder, f)
            if f.endswith((".xlsx", ".pptx", ".docx")):
                with zipfile.ZipFile(p) as z:
                    blob = b"".join(z.read(n) for n in z.namelist())
            elif f.endswith(".pdf"):
                blob = open(p, "rb").read()
            else:
                continue
            if any(b_ in blob for b_ in banned):
                leaky.append(f)
    A("realism: containers scrubbed, no writing library named in any office file or PDF", not leaky, leaky)
    stamp = B.EXPORT_STAMP.timestamp()
    A("realism: every pack file carries the pack's export time", all(abs(os.path.getmtime(os.path.join(target, f)) -
                                                                         stamp) < 1 for f in files))
    dates = []
    for f in files:
        if f.endswith(".xlsx"):
            with zipfile.ZipFile(os.path.join(target, f)) as z:
                core = z.read("docProps/core.xml").decode()
            dates += re.findall(r"(\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2})Z", core)
    A("realism: every workbook's saved stamps precede the export", dates and all(d <= "2026-10-09T19:42:00" for d in
                                                                                dates), max(dates))
    late = {}
    for f in files:
        if not f.endswith((".csv", ".txt", ".json")):
            continue
        txt = open(os.path.join(target, f), encoding="utf-8").read()
        ds = [m for m in re.findall(r"\b(20\d\d-\d\d-\d\d)\b", txt)]
        ds += ["%s-%s-%s" % (y, m_, d_) for m_, d_, y in re.findall(r"\b(\d\d)/(\d\d)/(20\d\d)\b", txt)
               if f == R.F_TA_OLD]
        ds += ["%s-%s-%s" % (y, m_, d_) for d_, m_, y in re.findall(r"\b(\d\d)/(\d\d)/(20\d\d)\b", txt)
               if f == R.F_HV["2025Q4"]]
        if ds and max(ds) > "2026-10-09":
            late[f] = max(ds)
    A("realism: no extract or note carries a date after the pack's export (9 October 2026)", not late, late)
    dash = []
    task = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    for f, t in T.items():
        if DASH.search(t):
            dash.append(f)
    for f in files:
        if f.endswith((".csv", ".txt", ".json")) and DASH.search(open(os.path.join(target, f),
                                                                                encoding="utf-8").read()):
            dash.append(f)
    for p in [os.path.join(task, "prompt.txt"), os.path.join(task, "DESIGN_NOTE.md"),
              os.path.join(task, "submission.md")] + [os.path.join(task, "generator", x) for x in
                                                       os.listdir(os.path.join(task, "generator")) if x.endswith(".py")]:
        if os.path.exists(p) and DASH.search(open(p, encoding="utf-8").read()):
            dash.append(os.path.basename(p))
    A("realism: no em or en dash in the pack, the prompt, the note, the submission or the generator", not dash, dash)
    people = [K.PEOPLE[k] for k in ("vp", "network", "treasury", "clearance", "hrsys")]
    A("realism: every Larchwood persona the documents name is the card's", all(any(p in t for t in T.values())
                                                                                 for p in people) and
      not any(K.PEOPLE["planner"] in t for t in T.values()))

    # ------------------------------------------------------------------ the golden files
    figs = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "figures.json"), encoding="utf-8"))
    A("golden: figures.json carries the call (Dorval 2, Calgary 1, held 3)",
      (figs["posts_placed"]["Dorval"], figs["posts_placed"]["Calgary"], figs["call"]["held"]) == ("2", "1", "3"))
    gf = sorted(os.listdir(golden))
    A("golden: two deliverables, the workbook and the deck", gf == sorted([DL.GOLD_BOOK, DL.GOLD_DECK]), gf)
    from pptx import Presentation
    prs = Presentation(os.path.join(golden, DL.GOLD_DECK))
    A("golden: the deck has two slides, the placement table and the overtime chart",
      len(prs.slides) == 2 and any(sh.has_table for sh in prs.slides[0].shapes) and
      any(sh.shape_type == 13 for sh in prs.slides[1].shapes))
    wbg = openpyxl.load_workbook(os.path.join(golden, DL.GOLD_BOOK), read_only=True)
    A("golden: the workbook's sheets are the placement, the forecast basis, the two ask tables and the notes",
      wbg.sheetnames == ["Placement", "Forecast basis", "Overtime FY2025-26", "Brandmoor FY2025-26", "Notes"],
      wbg.sheetnames)
    A.report()
    return A
