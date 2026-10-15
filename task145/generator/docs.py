"""The pack's documents: BRK-STD-03 (the brokerage staffing standard), CLR-OPS-11 (the clearance routing section of the
clearance operations manual), the 2027 and 2026 network plans, Brandmoor's rate schedule and statements, the HR staff
workbook, the AP register, the legacy gateway report and the two extracts notes."""
from __future__ import annotations

import datetime as dt

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import letter                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import cm                                                    # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import common as K   # noqa: E402
import field as FD   # noqa: E402
import records as R  # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/"
for _n, _f in (("Carlito", "crosextra/Carlito-Regular.ttf"), ("Carlito-Bold", "crosextra/Carlito-Bold.ttf"),
               ("Carlito-Italic", "crosextra/Carlito-Italic.ttf"), ("Caladea", "crosextra/Caladea-Regular.ttf"),
               ("Caladea-Bold", "crosextra/Caladea-Bold.ttf")):
    pdfmetrics.registerFont(TTFont(_n, FONT_DIR + _f))
SANS, SANS_B, SANS_I = "Carlito", "Carlito-Bold", "Carlito-Italic"
INK = colors.HexColor("#1d2229")
GREY = colors.HexColor("#5a6270")
RULE = colors.HexColor("#a3aab3")
PALE = colors.HexColor("#e8eef0")
BRAND = colors.HexColor("#1f5130")       # Larchwood green
BRAND2 = colors.HexColor("#5b2a86")      # Brandmoor

BODY = ParagraphStyle("body", fontName=SANS, fontSize=10.2, leading=13.6, textColor=INK, spaceAfter=5)
CLAUSE = ParagraphStyle("clause", parent=BODY, leftIndent=26, firstLineIndent=-26)
H1 = ParagraphStyle("h1", fontName=SANS_B, fontSize=15.0, leading=19, textColor=BRAND, spaceAfter=3)
H1S = ParagraphStyle("h1s", fontName=SANS, fontSize=9.4, leading=12.2, textColor=GREY, spaceAfter=10)
H2 = ParagraphStyle("h2", fontName=SANS_B, fontSize=11.2, leading=14, textColor=INK, spaceBefore=7, spaceAfter=3)
CELL = ParagraphStyle("cell", fontName=SANS, fontSize=8.8, leading=10.6, textColor=INK)
CELLR = ParagraphStyle("cellr", parent=CELL, alignment=2)
CELLB = ParagraphStyle("cellb", fontName=SANS_B, fontSize=8.8, leading=10.6, textColor=INK)
CELLBR = ParagraphStyle("cellbr", parent=CELLB, alignment=2)
SMALL = ParagraphStyle("small", fontName=SANS_I, fontSize=8.4, leading=10.8, textColor=GREY, spaceBefore=4)
B_H1 = ParagraphStyle("bh1", fontName="Caladea-Bold", fontSize=14.5, leading=18, textColor=BRAND2, spaceAfter=2)
B_BODY = ParagraphStyle("bbody", fontName="Caladea", fontSize=9.6, leading=12.4, textColor=INK, spaceAfter=3)


def _footer(left, right):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont(SANS, 7.8)
        canvas.setFillColor(GREY)
        canvas.drawString(2.0 * cm, 1.2 * cm, left)
        canvas.drawRightString(letter[0] - 2.0 * cm, 1.2 * cm, "%s  |  page %d" % (right, doc.page))
        canvas.restoreState()
    return draw


def _doc(path, title, author):
    return SimpleDocTemplate(path, pagesize=letter, leftMargin=2.1 * cm, rightMargin=2.1 * cm, topMargin=1.9 * cm,
                             bottomMargin=2.0 * cm, title=title, author=author, subject="", creator="", producer="")


def _fit(rows, total, size=8.2, pad=3.0, grow=(0, 1)):
    """Column widths for a billing table: each column as wide as its longest header word and its longest value, the
    room left given to the text columns."""
    need = []
    for j in range(len(rows[0])):
        w = 0.0
        for i, r in enumerate(rows):
            txt = str(r[j]).replace("&nbsp;", " ")
            parts = txt.split("<br/>") if i == 0 else [txt]
            for ptxt in parts:
                words = ptxt.split(" ") if i == 0 else [ptxt]
                for wd in words:
                    w = max(w, pdfmetrics.stringWidth(wd, SANS_B if i == 0 else SANS, size))
        need.append(w + 2 * pad + 1.0)
    spare = total - sum(need)
    assert spare >= 0, (spare, need)
    for j in grow:
        need[j] += spare / len(grow)
    return need


def _table(rows, widths, head=True, right=(), bold_last=False, size=None, pad=None):
    data = []
    cs = {}
    if size:
        for nm, base in (("c", CELL), ("cr", CELLR), ("b", CELLB), ("br", CELLBR)):
            cs[nm] = ParagraphStyle(nm + "s", parent=base, fontSize=size, leading=size * 1.22)
    for i, r in enumerate(rows):
        line = []
        for j, c in enumerate(r):
            hb = (head and i == 0) or (bold_last and i == len(rows) - 1)
            if size:
                st = (cs["br"] if j in right else cs["b"]) if hb else (cs["cr"] if j in right else cs["c"])
            else:
                st = (CELLBR if j in right else CELLB) if hb else (CELLR if j in right else CELL)
            line.append(Paragraph(str(c), st))
        data.append(line)
    t = Table(data, colWidths=widths, hAlign="LEFT")
    st = [("GRID", (0, 0), (-1, -1), 0.4, RULE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
          ("TOPPADDING", (0, 0), (-1, -1), 2.0), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.0)]
    if pad is not None:
        st += [("LEFTPADDING", (0, 0), (-1, -1), pad), ("RIGHTPADDING", (0, 0), (-1, -1), pad)]
    if head:
        st.append(("BACKGROUND", (0, 0), (-1, 0), PALE))
    t.setStyle(TableStyle(st))
    return t


def _c(n, text):
    return Paragraph("<b>%s</b>&nbsp;&nbsp;%s" % (n, text), CLAUSE)


def money(v):
    s = "{:,.2f}".format(abs(v))
    return ("-" + s) if v < 0 else s


def fr_money(v):
    s = "{:,.2f}".format(abs(v)).replace(",", "\u00a0").replace(".", ",")
    return ("-" + s if v < 0 else s) + "\u00a0$"


FR_MONTH = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre",
            "novembre", "décembre"]
EN_MONTH = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
            "November", "December"]


# ---------------------------------------------------------------------------------------------------- BRK-STD-03
def write_standard(path):
    P = K.PEOPLE
    est = K.ESTABLISHMENT
    s = [Paragraph("Larchwood Express Canada Ltd. &middot; Customs Brokerage", H1S),
         Paragraph("BRK-STD-03 Brokerage staffing standard", H1),
         Paragraph("Revision 4, effective 1 April 2025. Owner: %s, Vice-President, Trade and Brokerage Services. "
                   "Replaces revision 3 of 9 May 2022." % P["vp"], H1S),
         Paragraph("1. Purpose and scope", H2),
         _c("1.1", "This standard sets how the establishment of broker posts on the brokerage desks is reviewed each "
                   "year and where new broker posts are placed."),
         _c("1.2", "It covers the broker posts on the five gateway brokerage desks, Mississauga, Dorval, Richmond, "
                   "Calgary and Winnipeg: licensed customs brokers and broker trainees. Clearance operations staff, "
                   "who release shipments under the courier low-value program, are outside it."),
         Paragraph("2. Terms", H2),
         _c("2.1", "Formal entry: the accounting declaration a broker files for a shipment the clearance system "
                   "routes to the brokerage desk (CLR-OPS-11, section 3). One shipment routed to the desk is one "
                   "formal entry."),
         _c("2.2", "Plan year: the calendar year after the review."),
         _c("2.3", "Network plan: the volume plan Network Planning issues each September for the plan year."),
         Paragraph("3. Responsibilities", H2),
         _c("3.1", "Brokerage workforce planning prepares the review and files it with the establishment record."),
         _c("3.2", "The vice-president, Trade and Brokerage Services, presents it to the Canada leadership team, which "
                   "approves the establishment for the plan year."),
         Paragraph("4. The establishment review", H2),
         _c("4.1", "The review is held each October and places the new broker posts funded for the plan year."),
         _c("4.2", "A gateway's load is the average month of its forecast formal entries in the plan year (every "
                   "formal entry for a shipment released at the gateway in the plan year, whoever transmits it), "
                   "divided by its broker posts."),
         _c("4.3", "The forecast is made on the network plan for the plan year: the plan's shipments for the gateway "
                   "and the plan's planning assumptions."),
         _c("4.4", "A gateway's broker posts are the posts on its establishment at 1 January of the plan year, filled "
                   "or vacant (appendix A), and the posts placed in the review."),
         _c("4.5", "New posts are placed one at a time. Each goes to the gateway with the highest load, counting the "
                   "posts already placed in this review, while that load exceeds the standard load of 560 formal "
                   "entries a month. Posts not placed are held to the next review."),
         _c("4.6", "The review also reports, for each desk and each quarter of the last fiscal year, the overtime hours "
                   "worked at the desk and what the contract broker charged for the desk's entries."),
         Paragraph("5. New posts", H2),
         _c("5.1", "A new post is filled by a broker trainee, who files under the supervision of a licensed customs "
                   "broker on the desk for the first twelve months."),
         _c("5.2", "Recruitment for a placed post opens after the leadership team's approval, for a start in the "
                   "first quarter of the plan year."),
         KeepTogether([
             Paragraph("Appendix A. Establishment at 1 January 2027", H2),
             Paragraph("Maintained by Brokerage Workforce Planning; updated 18 September 2026 after the budget "
                       "approval of 16 September 2026.", SMALL),
             _table([["Desk", "Gateway", "Broker posts"]] +
                    [[K.DESK[g], g, str(est[g])] for g in K.GATEWAYS] + [["Total", "", str(sum(est.values()))]],
                    [5.0 * cm, 3.0 * cm, 3.4 * cm], right=(2,), bold_last=True),
             Paragraph("New broker posts funded for 2027: %d." % K.NEW_POSTS, SMALL)]),
         ]
    d = _doc(path, "BRK-STD-03 Brokerage staffing standard", P["vp"])
    f = _footer("Larchwood Express Canada Ltd., Customs Brokerage", "BRK-STD-03 rev. 4")
    d.build(s, onFirstPage=f, onLaterPages=f)


# ---------------------------------------------------------------------------------------------------- CLR-OPS-11
def _page_field(par):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        r = par.add_run()
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        elif text.strip() == "PAGE":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            r._r.append(it)
        else:
            r.text = text


def write_routing(path):
    """CLR-OPS-11, the clearance operations manual's valuation, routing and high-value check sections, as Clearance
    Operations keeps it in Word."""
    import docx
    from docx.shared import Cm, Pt, RGBColor
    P = K.PEOPLE
    doc = docx.Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(5)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
    fp = sec.footer.paragraphs[0]
    fp.text = "Larchwood Express Canada Ltd., Clearance Operations. CLR-OPS-11 rev. 7, page "
    _page_field(fp)
    for r in fp.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x5A, 0x62, 0x70)
    grey, green = RGBColor(0x5A, 0x62, 0x70), RGBColor(0x1F, 0x51, 0x30)

    def para(text, size=10.5, bold=False, color=None, after=5, before=0):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
        return p

    def clause(n, text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-1.0)
        r = p.add_run(n)
        r.bold = True
        p.add_run("\t" + text)
        return p

    para("Larchwood Express Canada Ltd. \u00b7 Clearance Operations", 9.5, color=grey, after=2)
    para("CLR-OPS-11 Clearance operations manual: valuation, routing and the high-value check", 15, True, green, 3)
    para("Revision 7, effective 21 October 2024. Owner: %s, Manager, Clearance Operations, Dorval." % P["clearance"],
         9.5, color=grey, after=10)
    para("1. Scope", 11.5, True, before=4, after=3)
    clause("1.1", "Inbound international shipments at the five import gateways: Mississauga (YYZ), Dorval (YUL), "
                  "Richmond (YVR), Calgary (YYC) and Winnipeg (YWG).")
    para("2. Valuation", 11.5, True, before=4, after=3)
    clause("2.1", "The clearance system values each shipment on its commercial invoice: the value for duty in the "
                  "invoice currency, as the shipper declares it and the invoice supports it.")
    clause("2.2", "A value in a foreign currency is converted to Canadian dollars at the customs exchange rate for the "
                  "day the shipment is released. The system holds the rate published for each business day; on other "
                  "days the last published rate applies.")
    para("3. Routing", 11.5, True, before=4, after=3)
    clause("3.1", "The system routes a shipment to the gateway's brokerage desk for a formal entry when (a) its value "
                  "for duty exceeds CAD 3,300, or (b) its goods fall under an other government department (OGD) "
                  "program, as the program flag on the tariff item shows, whatever its value.")
    clause("3.2", "A shipment invoiced in Canadian dollars is routed on its invoice value.")
    clause("3.3", "Every other shipment is released under the courier low-value program (CLVS) by clearance "
                  "operations.")
    clause("3.4", "A desk that cannot file an entry within its service window may send it to the contract broker, which "
                  "transmits it on Larchwood's account.")
    para("4. The high-value check", 11.5, True, before=4, after=3)
    clause("4.1", "Before release, clearance operations checks against its commercial invoice every shipment with a "
                  "value for duty of CAD 2,500 or more, and every OGD shipment.")
    clause("4.2", "The check's extract is pulled each calendar quarter: one row for each shipment released in the "
                  "quarter, and the shipments whose values were amended since the last pull.")
    para("5. Holds and examinations", 11.5, True, before=4, after=3)
    clause("5.1", "A shipment held for examination is released, and counted, on the day the hold is lifted.")
    clause("5.2", "A shipment refused or returned before release is not routed and does not appear in the check.")
    doc.core_properties.author = P["clearance"]
    doc.core_properties.title = "CLR-OPS-11 Valuation, routing and the high-value check"
    doc.save(path)


# ---------------------------------------------------------------------------------------------------- Brandmoor
def write_rate_schedule(path):
    f25, f26 = FD.FEES[2025], FD.FEES[2026]
    s = [Paragraph("Brandmoor Customs Services Inc.", B_H1),
         Paragraph("Licensed customs broker &middot; 4100 Matheson Boulevard East, Mississauga, Ontario", B_BODY),
         Spacer(1, 6),
         Paragraph("Schedule B. Contract rates for Larchwood Express Canada Ltd.", H2),
         Paragraph("Agreement LXC-2024-07 of 1 April 2024, schedule B as amended 12 January 2026.", B_BODY),
         _table([["Service", "Statements issued 1 April 2025 to 31 January 2026", "Statements issued from 1 February "
                  "2026"],
                 ["Formal entry transmitted, standard", "CAD %.2f per entry" % f25["std"], "CAD %.2f per entry" %
                  f26["std"]],
                 ["Formal entry transmitted, OGD program (permits and licences)", "CAD %.2f per entry" % f25["ogd"],
                  "CAD %.2f per entry" % f26["ogd"]],
                 ["Rush transmission (within two hours of the desk's request)", "CAD %.2f per entry, added" %
                  f25["rush"], "CAD %.2f per entry, added" % f26["rush"]],
                 ["Minimum monthly charge, per Larchwood desk with any entry on the statement", "CAD {:,.2f}".format(
                     f25["minimum"]), "CAD {:,.2f}".format(f26["minimum"])]],
                [6.6 * cm, 5.1 * cm, 5.1 * cm]),
         Spacer(1, 6),
         Paragraph("1. Brandmoor issues one statement a month for each Larchwood desk that sent entries, on the first "
                   "business day of the month, through the client portal and by email to Accounts Payable.", B_BODY),
         Paragraph("2. Where a desk's entry charges on a statement are below the minimum monthly charge, the "
                   "statement adds the difference as a minimum charge adjustment.", B_BODY),
         Paragraph("3. Charges are in Canadian dollars, before GST, HST or QST, which are added at the rate of the "
                   "desk's province.", B_BODY),
         Paragraph("4. A disputed statement is re-issued in full; the original is reversed by a credit memo.", B_BODY),
         ]
    d = _doc(path, "Brandmoor Customs Services, schedule B", "Brandmoor Customs Services Inc.")
    f = _footer("Brandmoor Customs Services Inc.", "Schedule B")
    d.build(s, onFirstPage=f, onLaterPages=f)


FRAME_W = letter[0] - 4.2 * cm - 2.0


def fr_rows(x, ch, gst, qst):
    return [["Relevé", "Déclarations<br/>ordinaires", "Déclarations<br/>MAG", "Urgentes", "Frais",
             "Rajustement<br/>au minimum", "Sous-total", "TPS", "TVQ", "Total"],
            [x.number, str(x.std), str(x.ogd), str(x.rush), fr_money(ch), fr_money(x.minimum_adj), fr_money(x.net),
             fr_money(gst), fr_money(qst), fr_money(x.total)]]


def re_rows(st, ch):
    return [["Statement", "Desk", "Standard<br/>entries", "OGD<br/>entries", "Rush<br/>requests", "Entry<br/>charges",
             "Minimum<br/>adjustment", "Subtotal", "Tax", "Total"],
            [st.number, "%s (%s)" % (K.DESK[st.gw], FD.CLIENT_CODE[st.gw]), str(st.std), str(st.ogd), str(st.rush),
             money(ch), money(st.minimum_adj), money(st.net),
             "%s %s" % ("/".join(n for n, _ in st.taxes), money(sum(v for _, v in st.taxes))), money(st.total)]]


def write_statements(path, stmts):
    """Brandmoor's statements as Accounts Payable received them, statement months April 2025 to March 2026: one page
    a month for the four desks billed from Mississauga, Dorval's statement in French from the Montreal office, and the
    re-issued statement."""
    by_issue = {}
    for st in stmts:
        if st.reissue_of:
            continue
        key = (st.issued.year, st.issued.month)
        if dt.date(2025, 4, 1) <= st.issued <= dt.date(2026, 3, 31):
            by_issue.setdefault(key, []).append(st)
    s = []
    first = True
    for key in sorted(by_issue):
        sts = sorted(by_issue[key], key=lambda x: K.GATEWAYS.index(x.gw))
        issued = sts[0].issued
        if not first:
            s.append(PageBreak())
        first = False
        en = [x for x in sts if x.gw != "YUL"]
        s += [Paragraph("Brandmoor Customs Services Inc.", B_H1),
              Paragraph("Statement of account &middot; Client: Larchwood Express Canada Ltd. &middot; Statement month: "
                        "%s %d &middot; Issued %d %s %d" % (EN_MONTH[key[1] - 1], key[0], issued.day,
                                                           EN_MONTH[issued.month - 1], issued.year), B_BODY),
              Spacer(1, 4)]
        rows = [["Statement", "Desk", "Standard<br/>entries", "OGD<br/>entries", "Rush<br/>requests", "Entry<br/>charges",
                 "Minimum<br/>adjustment", "Subtotal", "Tax", "Total"]]
        for x in en:
            ch = round(x.std * x.std_rate + x.ogd * x.ogd_rate + x.rush * x.rush_rate, 2)
            rows.append([x.number, "%s (%s)" % (K.DESK[x.gw], FD.CLIENT_CODE[x.gw]), str(x.std), str(x.ogd),
                         str(x.rush), money(ch), money(x.minimum_adj), money(x.net),
                         "%s %s" % ("/".join(n for n, _ in x.taxes), money(sum(v for _, v in x.taxes))),
                         money(x.total)])
        s.append(_table(rows, _fit(rows, FRAME_W), right=(2, 3, 4, 5, 6, 7, 8, 9), size=8.2, pad=3.0))
        r0 = sts[0]
        s.append(Paragraph("Rates: standard CAD %.2f, OGD CAD %.2f, rush CAD %.2f per entry; minimum monthly charge "
                           "CAD {:,.2f} per desk (schedule B). Payment terms 30 days.".format(
                               FD.FEES[2026 if issued >= dt.date(2026, 2, 1) else 2025]["minimum"]) %
                           (r0.std_rate, r0.ogd_rate, r0.rush_rate), SMALL))
        fr = [x for x in sts if x.gw == "YUL"]
        if fr:
            x = fr[0]
            ch = round(x.std * x.std_rate + x.ogd * x.ogd_rate + x.rush * x.rush_rate, 2)
            gst = [v for n, v in x.taxes if n == "GST"][0]
            qst = [v for n, v in x.taxes if n == "QST"][0]
            s += [Spacer(1, 12), Paragraph("Brandmoor Services en douane inc. &middot; Bureau de Montréal", B_H1),
                  Paragraph("Relevé de compte &middot; Client : Larchwood Express Canada Ltd., comptoir de Dorval (%s) "
                            "&middot; Mois du relevé : %s %d &middot; Émis le %d %s %d" % (
                                FD.CLIENT_CODE["YUL"], FR_MONTH[key[1] - 1], key[0], issued.day,
                                FR_MONTH[issued.month - 1], issued.year), B_BODY),
                  _table(fr_rows(x, ch, gst, qst), _fit(fr_rows(x, ch, gst, qst), FRAME_W, grow=(0,)),
                         right=(1, 2, 3, 4, 5, 6, 7, 8, 9), size=8.2, pad=3.0),
                  Paragraph("MAG : déclarations visées par un programme d'un autre ministère. Taux selon l'annexe B.",
                            SMALL)]
    for st in stmts:
        if not st.reissue_of:
            continue
        ch = round(st.std * st.std_rate + st.ogd * st.ogd_rate + st.rush * st.rush_rate, 2)
        s += [PageBreak(), Paragraph("Brandmoor Customs Services Inc.", B_H1),
              Paragraph("Re-issued statement &middot; Client: Larchwood Express Canada Ltd. &middot; Statement month: "
                        "%s %d &middot; Re-issued %d %s %d &middot; Replaces statement %s" % (
                            EN_MONTH[(st.month[1] % 12)], st.month[0] + (st.month[1] == 12), st.issued.day,
                            EN_MONTH[st.issued.month - 1], st.issued.year, st.reissue_of), B_BODY),
              _table(re_rows(st, ch), _fit(re_rows(st, ch), FRAME_W), right=(2, 3, 4, 5, 6, 7, 8, 9), size=8.2,
                     pad=3.0),
              Paragraph("Re-issued on the client's query of 9 March 2026: rush transmissions corrected to the desk's "
                        "requests. A credit memo reversing statement %s follows." % st.reissue_of, SMALL)]
    d = _doc(path, "Brandmoor statements of account", "Brandmoor Customs Services Inc.")
    f = _footer("Brandmoor Customs Services Inc., statements of account", "Larchwood Express Canada Ltd.")
    d.build(s, onFirstPage=f, onLaterPages=f)


# ---------------------------------------------------------------------------------------------------- workbooks
def _wb():
    from openpyxl import Workbook
    wb = Workbook()
    return wb


def _style_head(ws, row, ncol):
    from openpyxl.styles import Alignment, Font, PatternFill
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name="Calibri", bold=True, size=10)
        cell.fill = PatternFill("solid", fgColor="DDE6E9")
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_plan(path, W, year):
    """The network plan for a plan year: its cover, assumptions, inbound volumes by gateway and month (thousands),
    outbound volumes and linehaul."""
    from openpyxl.styles import Font
    if year == 2027:
        base_months, growth, issued = K.BASE_MONTHS, K.PLAN_GROWTH, K.PLAN_ISSUED
        rates = K.PLAN_RATE
        rate_date = "4 September 2026"
    else:
        base_months = K.months(dt.date(2024, 10, 1), dt.date(2025, 9, 30))
        growth = {"YYZ": 0.031, "YUL": 0.036, "YVR": 0.044, "YYC": 0.022, "YWG": 0.027}
        issued = dt.date(2025, 9, 19)
        rates = {"USD": 1.3900, "EUR": 1.5800, "GBP": 1.8400, "CNY": 0.19200, "JPY": 0.0091500}
        rate_date = "5 September 2025"
    by = base_months[0][0]
    wb = _wb()
    ws = wb.active
    ws.title = "Cover"
    lines = ["Larchwood Express Canada Ltd.", "Network plan %d" % year,
             "Issued %d %s %d by Network Planning (%s, Manager, Network Planning)" % (
                 issued.day, EN_MONTH[issued.month - 1], issued.year, K.PEOPLE["network"]),
             "Planning rates from Treasury (%s), %s" % (K.PEOPLE["treasury"], rate_date),
             "", "Sheets: Assumptions, Inbound volumes, Outbound volumes, Linehaul",
             "Volumes are thousands of shipments; the plan year is January to December %d." % year]
    for i, t in enumerate(lines, 1):
        if t:
            ws.cell(row=i, column=1, value=t)
    ws.cell(row=1, column=1).font = Font(bold=True, size=13)
    ws.cell(row=2, column=1).font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 92
    # assumptions
    a = wb.create_sheet("Assumptions")
    bm0, bm1 = base_months[0], base_months[-1]
    base_label = "%s %d to %s %d" % (EN_MONTH[bm0[1] - 1], bm0[0], EN_MONTH[bm1[1] - 1], bm1[0])
    gtxt = ", ".join("%s %+.1f%%" % (K.PLAN_REGION[g], 100 * growth[g]) for g in K.GATEWAYS)
    rtxt = ", ".join(("%s %." + str(K.RATE_DP[c]) + "f") % (c, rates[c]) for c in ["USD", "EUR", "GBP", "CNY", "JPY"])
    rows = [("1", "Base", "The twelve months %s." % base_label),
            ("2", "Shipment profile", "Each gateway's lanes, customers and goods as in the base."),
            ("3", "Inbound volume growth on the base", gtxt + "."),
            ("4", "Outbound volume growth on the base", "Network %+.1f%%, by gateway on the Outbound volumes sheet." %
             (100 * (sum(growth.values()) / 5 - 0.004))),
            ("5", "Linehaul", "The %d schedule on the Linehaul sheet; no new gateway." % year),
            ("6", "Fuel", "Jet fuel CAD %.3f per litre, road diesel CAD %.3f per litre." % (
                (0.982, 1.648) if year == 2027 else (1.034, 1.702))),
            ("7", "Exchange rates", "Treasury planning rates for %d, CAD per unit: %s." % (year, rtxt)),
            ("8", "Carrier pricing", "General rate increase %.1f%% from %s." % (
                (5.4, "4 January 2027") if year == 2027 else (5.9, "5 January 2026"))),
            ("9", "Labour", "Wage settlement %.1f%% from 1 April %d; establishment as approved." % (
                3.1 if year == 2027 else 3.4, year))]
    a.append(["No.", "Assumption", "Plan %d" % year])
    _style_head(a, 1, 3)
    for r_ in rows:
        a.append(list(r_))
    a.column_dimensions["A"].width = 6
    a.column_dimensions["B"].width = 34
    a.column_dimensions["C"].width = 120
    # inbound volumes
    v = wb.create_sheet("Inbound volumes")
    v.append(["Inbound international shipments, thousands"])
    v.append(["Region (gateway)"] + ["%s %d" % (EN_MONTH[m - 1][:3], year) for m in range(1, 13)] +
             ["Plan %d" % year, "Base %s" % base_label, "Growth"])
    _style_head(v, 2, 16)
    for g in K.GATEWAYS:
        base = [W.volume[(g, y, m)] for (y, m) in base_months] if (base_months[0] in [(y, m) for (y, m) in
                                                                                      K.HIST_MONTHS]) else None
        bymonth = {m: W.volume[(g, y, m)] for (y, m) in base_months}
        plan_k = [K.round_dp(bymonth[m] * (1 + growth[g]) / 1000.0, 1) for m in range(1, 13)]
        base_k = K.round_dp(sum(base) / 1000.0, 1)
        v.append([K.PLAN_REGION[g]] + plan_k + [round(sum(plan_k), 1), base_k, "%+.1f%%" % (100 * growth[g])])
    tot = [round(sum(v.cell(row=r_, column=c).value for r_ in range(3, 8)), 1) for c in range(2, 16)]
    v.append(["Network"] + tot + ["%+.1f%%" % (100 * (tot[12] / tot[13] - 1))])
    v.column_dimensions["A"].width = 34
    # outbound and linehaul
    o = wb.create_sheet("Outbound volumes")
    o.append(["Outbound international shipments, thousands"])
    o.append(["Region (gateway)", "Plan %d" % year, "Base", "Growth"])
    _style_head(o, 2, 4)
    share = {"YYZ": 0.46, "YUL": 0.38, "YVR": 0.29, "YYC": 0.34, "YWG": 0.27}
    for g in K.GATEWAYS:
        base_o = sum(W.volume[(g, y, m)] for (y, m) in base_months) * share[g] / 1000.0
        go = growth[g] - 0.004
        o.append([K.PLAN_REGION[g], round(base_o * (1 + go), 1), round(base_o, 1), "%+.1f%%" % (100 * go)])
    o.column_dimensions["A"].width = 34
    lh = wb.create_sheet("Linehaul")
    lh.append(["Scheduled inbound international rotations a week, %d" % year])
    lh.append(["Gateway", "Aircraft", "Rotations a week", "Origin hubs"])
    _style_head(lh, 2, 4)
    sched = {"YYZ": ("B767-300F", 11, "Cincinnati, Leipzig, East Midlands"), "YUL": ("B767-300F", 6, "Leipzig, Cincinnati"),
             "YVR": ("B767-300F", 7, "Hong Kong, Incheon, Anchorage"), "YYC": ("B737-800BCF", 5, "Cincinnati"),
             "YWG": ("B737-800BCF", 4, "Cincinnati, Minneapolis")}
    for g in K.GATEWAYS:
        lh.append([K.DESK[g] + " (" + g + ")"] + list(sched[g]))
    lh.column_dimensions["A"].width = 22
    lh.column_dimensions["D"].width = 40
    wb.save(path)


def write_gw_legacy(path, rows):
    from openpyxl.styles import Font
    wb = _wb()
    ws = wb.active
    ws.title = "Gateway activity"
    ws.append(["Larchwood Express Canada Ltd. - Customs Brokerage"])
    ws.append(["Gateway activity, monthly, January 2021 to September 2024 (brokerage system BX4 report GA-02)"])
    ws.append(["Formal entries are counted in the month their B3 accounting was accepted."])
    ws.append([])
    ws.append(["Month", "Gateway", "Inbound international shipments", "Released under CLVS",
               "Formal entries (B3 accounted)"])
    _style_head(ws, 5, 5)
    ws.cell(row=1, column=1).font = Font(bold=True, size=12)
    for r_ in rows:
        ws.append(r_)
        ws.cell(row=ws.max_row, column=1).number_format = "mmm-yy"
    for c, w in zip("ABCDE", (10, 14, 18, 16, 18)):
        ws.column_dimensions[c].width = w
    wb.save(path)


def write_hr_staff(path, staff):
    wb = _wb()
    ws = wb.active
    ws.title = "Brokerage staff"
    ws.append(["Employee number", "Legacy ID", "Surname", "Given name", "Position", "Customs broker licence", "Desk",
               "Hire date", "Status", "Termination date"])
    _style_head(ws, 1, 10)
    for s in sorted(staff, key=lambda x: (K.GATEWAYS.index(x.desk_on(FD.HR_PULL) if (x.left is None) else
                                                                x.desk_on(x.left)), x.last)):
        desk = s.desk_on(FD.HR_PULL) if s.left is None else s.desk_on(s.left)
        ws.append([s.emp_no, s.legacy_id, s.last, s.first, s.position, s.licence, K.DESK[desk], s.hired,
                   "Active" if s.left is None else "Terminated", s.left])
        ws.cell(row=ws.max_row, column=8).number_format = "yyyy-mm-dd"
        if s.left:
            ws.cell(row=ws.max_row, column=10).number_format = "yyyy-mm-dd"
    for c, w in zip("ABCDEFGHIJ", (16, 11, 16, 14, 24, 20, 13, 12, 11, 15)):
        ws.column_dimensions[c].width = w
    wb.save(path)


def write_ap(path, rows):
    wb = _wb()
    ws = wb.active
    ws.title = "AP register"
    ws.append(["Supplier", "Document type", "Document number", "Document date", "Posting date", "Cost centre",
               "Net CAD", "Tax CAD", "Gross CAD", "Reference"])
    _style_head(ws, 1, 10)
    for r_ in rows:
        ws.append(r_)
        for c in (4, 5):
            ws.cell(row=ws.max_row, column=c).number_format = "yyyy-mm-dd"
        for c in (7, 8, 9):
            ws.cell(row=ws.max_row, column=c).number_format = "#,##0.00"
    for c, w in zip("ABCDEFGHIJ", (30, 14, 20, 13, 13, 11, 12, 11, 12, 28)):
        ws.column_dimensions[c].width = w
    wb.save(path)


# ---------------------------------------------------------------------------------------------------- notes
def write_clearance_note(path):
    t = """Clearance and brokerage extracts for the 2027 establishment review
Pulled 7 October 2026 by Clearance Operations for Customs Brokerage (workforce planning).

Files
  highvalue_check_2024Q4.csv to highvalue_check_2026Q3.csv (eight files)
      The high-value check extract for each calendar quarter (CLR-OPS-11, 4.2), October 2024 to September 2026; the
      four files to 2025Q3 are the extracts pulled for the 2026 review. The clearance system release of 17 February
      2026 changed the extract's layout; the files to 2025Q4, the last pulled 6 January 2026, are in the earlier layout.
  gateway_activity_2021-01_2024-09.xlsx
      The legacy brokerage system's monthly gateway report (GA-02), to its last month, September 2024.
  gateway_activity_2024-10_2026-09.csv
      The clearance system's monthly gateway report, from October 2024.
  export_declarations_monthly_2025-04_2026-09.csv
      The export desk's monthly report by gateway: export declarations filed, CERS exemptions and outbound
      international shipments, from April 2025.

High-value check (earlier layout field / current layout field)
  WAYBILL / waybill                         the shipment's waybill number, ten digits with a check digit
  GTWY / gateway                            the import gateway: YYZ Mississauga, YUL Dorval, YVR Richmond,
                                            YYC Calgary, YWG Winnipeg
  REL_DT / release_local_date               the day of release at the gateway, local time (REL_DT is DD/MM/YYYY)
  REL_TM_UTC / release_utc                  the time of release in UTC
  IMPORTER_ACCT / importer_account          the importer's business number and import account
  HS_CODE / hs_code                         the tariff classification, ten digits
  ORIG_CTRY / origin_country                country of origin, ISO 3166 alpha-2
  INV_CURR / invoice_currency               the currency of the commercial invoice
  VFD_INV / value_for_duty                  the value for duty in the invoice currency
  CUST_RATE / customs_exchange_rate         CAD per unit of the invoice currency, the customs rate of the release day
  VFD_CAD / value_for_duty_cad              the value for duty in CAD
  ROUTE / route                             FORMAL: routed to the brokerage desk for a formal entry;
                                            CLVS: released under the courier low-value program
  ENTRY_NO / entry_number                   the formal entry's transaction number (FORMAL rows)
  OGD_PGM / ogd_program                     the OGD program on the tariff item (CFIA, HC-COS, HC-DRUG, HC-NHP, HC-PMRA,
                                            HC-CPS, ISED-RE); blank in the earlier layout and NONE in the current one
                                            where none applies
  AMEND / amended                           Y on a row re-extracted because its value was amended after the previous
                                            pull; the row carries the amended value
One row per shipment in each extract; the extract covers the quarter's releases by local release date and the
shipments amended since the previous pull.

Gateway reports
  Legacy (GA-02): Month, Gateway (desk name), Inbound international shipments, Released under CLVS, Formal entries
  (B3 accounted). Formal entries are counted in the month their B3 accounting was accepted, a few business days after
  release.
  Current: gateway_code, period (YYYY-MM), inbound_international_shipments, clvs_releases, formal_entries_released.
  Formal entries are counted in the month of release; clvs_releases are the month's other inbound international
  shipments.
"""
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(t)


def write_payroll_note(path):
    t = """Payroll, HR and contract broker extracts for the 2027 establishment review
Pulled 6 October 2026 by HR Systems (%s) for Customs Brokerage (workforce planning).
Larchwood's fiscal year runs April to March; fiscal 2025-26 is 1 April 2025 to 31 March 2026.

Files
  ta_legacy_export_2025-04-01_2025-10-31.csv      legacy time and attendance system, final export of 31 October 2025
  ta_export_2025-10-27_2026-09-30.csv             current time and attendance system, live from 1 November 2025
  hr_brokerage_staff_2026-10-06.xlsx              the brokerage desks' staff at the pull, with terminations
                                                  since 1 April 2025
  hr_assignment_changes_2025-04_2026-09.csv       hires, transfers, terminations and licence grants
  finance_cost_centres_2026.csv                   cost centres in the legacy and the ERP 2025 schemes
  brandmoor_portal_detail_2025-04_2026-03.json    Brandmoor's client portal statement detail, statement months April
                                                  2025 to March 2026
  brandmoor_statements_2025-04_2026-03.pdf        Brandmoor's statements as received, April 2025 to March 2026
  ap_invoice_register_brandmoor_2025-04_2026-09.xlsx  Accounts Payable postings for Brandmoor, April 2025 to
                                                  September 2026
  brandmoor_rate_schedule_2025.pdf                schedule B of the Brandmoor agreement

Time and attendance, legacy export
  EMPLID      legacy employee id
  WORK_DATE   the day worked, MM/DD/YYYY
  PAY_CODE    REG regular, OT15 overtime at time and a half, OT20 overtime at double time, CALLB call-back
              overtime, VAC vacation, SICK sick leave, TRAIN training, STAT statutory holiday
  HOURS       the hours the line pays
  IN_TIME, OUT_TIME   punches, local time
  HOME_CC     the employee's home cost centre (legacy scheme)

Time and attendance, current export
  employee_number   employee number
  work_date         the day worked
  earning_code      REG, OT1.5, OT2.0, CB call-back overtime, VAC, SCK, TRN, STH statutory holiday
  quantity          the hours the line pays
  punch_in, punch_out   punches, local time, rounded to the quarter hour
  home_cost_center  the employee's home cost centre (ERP 2025 scheme)
  source            TIMECLOCK, MANUAL (entered by a supervisor) or MIGRATED (loaded at the system's start)

HR assignment changes
  employee_number, effective_date (the first day under the change), change_type (HIRE, TRANSFER, TERMINATION,
  LICENCE), from_desk, to_desk, note

Brandmoor portal detail (the lines array)
  statement_month, statement_number, client_code (LXC-0411 Mississauga, LXC-0412 Dorval, LXC-0415 Richmond,
  LXC-0418 Calgary, LXC-0419 Winnipeg), entry_number, transmitted_et (Eastern time), service (STD standard,
  OGD an entry under an OGD program), rush (Y), fee, rush_fee, statement_minimum_adjustment
""" % K.PEOPLE["hrsys"]
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(t)
