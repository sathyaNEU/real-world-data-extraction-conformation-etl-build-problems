"""The golden deliverables (broker_posts_2027.xlsx, establishment_review_2027.pptx) and generator/figures.json, from one
computation of the golden: the placement on the forecast at the plan's rates, and the two asks."""
from __future__ import annotations

import io
import json
import os

import common as K
import calc as C

GOLD_BOOK = "broker_posts_2027.xlsx"
GOLD_DECK = "establishment_review_2027.pptx"
FQ_LABEL = {1: "Q1 Apr-Jun 2025", 2: "Q2 Jul-Sep 2025", 3: "Q3 Oct-Dec 2025", 4: "Q4 Jan-Mar 2026"}
SERIES = {"YYZ": "#2a78d6", "YUL": "#eb6834", "YVR": "#1baf7a", "YYC": "#eda100", "YWG": "#e87ba4"}


def golden(W, staff, lines, stmts):
    E = C.golden_E(W)
    placed, held, final, order = C.placement(E)
    Ed = C.default_E(W)
    pd_, hd, fd, od = C.placement(Ed)
    A = C.ask_a(staff, lines, "golden")
    B = C.ask_b(stmts, "golden")
    n = {g: C.formal_at(W, g, lambda s: K.PLAN_RATE[s.cur]) for g in K.GATEWAYS}
    return dict(E=E, placed=placed, held=held, final=final, order=order, Ed=Ed, placed_d=pd_, held_d=hd, final_d=fd,
                A=A, B=B, n=n, ships=len(W.ships))


def figures(G):
    d = K.DESK
    f = {"call": {"posts_dorval": str(G["placed"]["YUL"]), "posts_calgary": str(G["placed"]["YYC"]),
                  "held": str(G["held"]), "new_posts": str(K.NEW_POSTS), "standard_load": str(K.STANDARD_LOAD),
                  "plan_year": "2027"},
         "posts_placed": {d[g]: str(G["placed"][g]) for g in K.GATEWAYS},
         "establishment": {d[g]: str(K.ESTABLISHMENT[g]) for g in K.GATEWAYS},
         "posts_after": {d[g]: str(K.ESTABLISHMENT[g] + G["placed"][g]) for g in K.GATEWAYS},
         "forecast_monthly_formal_entries": {d[g]: "{:,}".format(K.round_to(G["E"][g], 50)) for g in K.GATEWAYS},
         "entries_per_post_after": {d[g]: str(K.round_to(G["final"][g], 5)) for g in K.GATEWAYS},
         "blocking_load": {"desk": d[max(G["final"], key=G["final"].get)],
                           "load": str(K.round_to(max(G["final"].values()), 5))},
         "plan_rates": {c: ("%." + str(K.RATE_DP[c]) + "f") % K.PLAN_RATE[c]
                        for c in ("USD", "EUR", "GBP", "CNY", "JPY")},
         "line_cad": "3,300",
         "fx_effect_pct": {d[g]: "%+.0f" % (100 * (G["E"][g] / G["Ed"][g] - 1)) for g in K.GATEWAYS},
         "default_path": {"posts_placed": {d[g]: str(G["placed_d"][g]) for g in K.GATEWAYS},
                          "held": str(G["held_d"]),
                          "forecast_monthly_formal_entries": {d[g]: "{:,}".format(K.round_to(G["Ed"][g], 50))
                                                              for g in K.GATEWAYS}},
         "ask_a_overtime_hours": {d[g]: {"FQ%d" % q: "{:,}".format(K.round_to(G["A"][(g, q)], 1)) for q in (1, 2, 3, 4)}
                                  for g in K.GATEWAYS},
         "ask_b_brandmoor_cad": {d[g]: {"FQ%d" % q: "{:,}".format(K.round_to(G["B"][(g, q)], 1)) for q in (1, 2, 3, 4)}
                                 for g in K.GATEWAYS},
         "fiscal_year": "2025-26",
         "base_year_shipments_checked": "{:,}".format(G["ships"]),
         "base_year_formal_at_plan_rates": {d[g]: "{:,}".format(G["n"][g]) for g in K.GATEWAYS},
         "placement_steps": ["%s %d" % (d[g], K.round_to(l, 5)) for g, l in G["order"]],
         "ask_a_total_hours": "{:,}".format(K.round_to(sum(G["A"].values()), 1)),
         "ask_b_total_cad": "{:,}".format(K.round_to(sum(G["B"].values()), 1))}
    return f


# ---------------------------------------------------------------------------------------------------- the workbook
def write_book(path, G):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    thin = Side(style="thin", color="A3AAB3")
    head_fill = PatternFill("solid", fgColor="E3EBE5")
    bold = Font(name="Calibri", bold=True, size=10)
    norm = Font(name="Calibri", size=10)
    title_f = Font(name="Calibri", bold=True, size=13, color="1F5130")

    def header(ws, row, cols):
        for c, v in enumerate(cols, 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = bold
            cell.fill = head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=thin)

    ws = wb.active
    ws.title = "Placement"
    ws["A1"] = "Two of the six new broker posts go to Dorval and one to Calgary; three are held to the 2028 review"
    ws["A1"].font = title_f
    ws["A2"] = ("Establishment review for 2027 under BRK-STD-03 section 4. Forecast on the 2027 network plan "
                "(issued 18 September 2026), its shipments and planning rates; standard load 560 formal entries a "
                "month per broker post.")
    ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="5A6270")
    cols = ["Desk", "Gateway", "Broker posts at 1 January 2027", "New posts placed", "Broker posts after the review",
            "Forecast 2027 formal entries, average month", "Formal entries a month per broker post after the review"]
    header(ws, 4, cols)
    r = 5
    for g in K.GATEWAYS:
        vals = [K.DESK[g], g, K.ESTABLISHMENT[g], G["placed"][g], K.ESTABLISHMENT[g] + G["placed"][g],
                K.round_to(G["E"][g], 50), K.round_to(G["final"][g], 5)]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = norm
            if c >= 3:
                cell.number_format = "#,##0"
        r += 1
    tot = [("Total", ""), sum(K.ESTABLISHMENT.values()), sum(G["placed"].values()),
           sum(K.ESTABLISHMENT.values()) + sum(G["placed"].values()), K.round_to(sum(G["E"].values()), 50)]
    ws.cell(row=r, column=1, value="All desks").font = bold
    for c, v in zip((3, 4, 5, 6), tot[1:]):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = bold
        cell.number_format = "#,##0"
        cell.border = Border(top=thin)
    r += 2
    ws.cell(row=r, column=1, value="New posts held to the 2028 review").font = bold
    ws.cell(row=r, column=4, value=G["held"]).font = bold
    r += 2
    ws.cell(row=r, column=1, value="Order of placement").font = bold
    r += 1
    for i, (g, _) in enumerate(G["order"], 1):
        ws.cell(row=r, column=1, value="Post %d" % i).font = norm
        ws.cell(row=r, column=2, value=K.DESK[g]).font = norm
        r += 1
    ws.cell(row=r, column=1, value="Fourth post").font = norm
    ws.cell(row=r, column=2, value="held: the highest load after three posts, %s, is under 560" %
            K.DESK[max(G["final"], key=G["final"].get)]).font = norm
    widths = (14, 10, 14, 11, 13, 17, 19)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.row_dimensions[4].height = 44
    ws.freeze_panes = "A5"

    # overtime
    ot = wb.create_sheet("Overtime FY2025-26")
    ot["A1"] = "Overtime hours worked at each desk, fiscal 2025-26"
    ot["A1"].font = title_f
    ot["A2"] = ("Overtime at time and a half and double time, and call-backs at the time worked, on the desk each "
                "broker was assigned to on the day; fiscal quarters April to March.")
    ot["A2"].font = Font(name="Calibri", size=9, italic=True, color="5A6270")
    header(ot, 4, ["Desk"] + [FQ_LABEL[q] for q in (1, 2, 3, 4)] + ["Fiscal 2025-26"])
    r = 5
    for g in K.GATEWAYS:
        ot.cell(row=r, column=1, value=K.DESK[g]).font = norm
        for q in (1, 2, 3, 4):
            c = ot.cell(row=r, column=1 + q, value=K.round_to(G["A"][(g, q)], 1))
            c.number_format = "#,##0"
            c.font = norm
        c = ot.cell(row=r, column=6, value=K.round_to(sum(G["A"][(g, q)] for q in (1, 2, 3, 4)), 1))
        c.number_format = "#,##0"
        c.font = norm
        r += 1
    ot.cell(row=r, column=1, value="All desks").font = bold
    for q in (1, 2, 3, 4):
        c = ot.cell(row=r, column=1 + q, value=K.round_to(sum(G["A"][(g, q)] for g in K.GATEWAYS), 1))
        c.number_format = "#,##0"
        c.font = bold
        c.border = Border(top=thin)
    c = ot.cell(row=r, column=6, value=K.round_to(sum(G["A"].values()), 1))
    c.number_format = "#,##0"
    c.font = bold
    c.border = Border(top=thin)
    for i, w in enumerate((14, 16, 16, 16, 16, 15)):
        ot.column_dimensions[chr(65 + i)].width = w
    ot.freeze_panes = "B5"

    # Brandmoor
    bm = wb.create_sheet("Brandmoor FY2025-26")
    bm["A1"] = "Brandmoor charges for the entries it transmitted for each desk, fiscal 2025-26"
    bm["A1"].font = title_f
    bm["A2"] = ("CAD before GST, HST and QST, by the quarter the entries were transmitted; entry charges and minimum "
                "charge adjustments as the statements bill them, the re-issued statement in place of the original.")
    bm["A2"].font = Font(name="Calibri", size=9, italic=True, color="5A6270")
    header(bm, 4, ["Desk"] + [FQ_LABEL[q] for q in (1, 2, 3, 4)] + ["Fiscal 2025-26"])
    r = 5
    for g in K.GATEWAYS:
        bm.cell(row=r, column=1, value=K.DESK[g]).font = norm
        for q in (1, 2, 3, 4):
            c = bm.cell(row=r, column=1 + q, value=K.round_to(G["B"][(g, q)], 1))
            c.number_format = "$#,##0"
            c.font = norm
        c = bm.cell(row=r, column=6, value=K.round_to(sum(G["B"][(g, q)] for q in (1, 2, 3, 4)), 1))
        c.number_format = "$#,##0"
        c.font = norm
        r += 1
    bm.cell(row=r, column=1, value="All desks").font = bold
    for q in (1, 2, 3, 4):
        c = bm.cell(row=r, column=1 + q, value=K.round_to(sum(G["B"][(g, q)] for g in K.GATEWAYS), 1))
        c.number_format = "$#,##0"
        c.font = bold
        c.border = Border(top=thin)
    c = bm.cell(row=r, column=6, value=K.round_to(sum(G["B"].values()), 1))
    c.number_format = "$#,##0"
    c.font = bold
    c.border = Border(top=thin)
    for i, w in enumerate((14, 16, 16, 16, 16, 15)):
        bm.column_dimensions[chr(65 + i)].width = w
    bm.freeze_panes = "B5"

    nt = wb.create_sheet("Notes")
    notes = [
        "Notes",
        "Load: a desk's forecast average monthly formal entries in 2027 over its broker posts, the posts already "
        "placed in the review counted (BRK-STD-03 4.2 to 4.5). A post is placed while the highest load exceeds 560.",
        "Forecast: each gateway's shipments of October 2025 to September 2026 in the high-value check, routed as "
        "CLR-OPS-11 section 3 routes them at the 2027 planning rates (USD 1.3400, EUR 1.6800, GBP 1.8750, CNY "
        "0.18800, JPY 0.0092500), scaled to the 2027 network plan's inbound volumes.",
        "Broker posts at 1 January 2027: BRK-STD-03 appendix A.",
        "Fiscal 2025-26 runs 1 April 2025 to 31 March 2026; Q1 is April to June 2025 and Q4 January to March 2026.",
        "Overtime: the legacy export to 31 October 2025 and the current system from 1 November (its migrated lines "
        "for 27 to 31 October are the legacy timecards, counted once); call-backs at the time between their punches, "
        "not the three-hour minimum they pay; each line on the desk the HR assignment changes give for the work date.",
        "Brandmoor: each statement on the month of the transmissions it bills, from the statements, the client "
        "portal detail and the AP register to 30 September 2026; March 2026 transmissions from the statements of "
        "1 April 2026. Mississauga's February 2026 transmissions as re-issued on 23 March 2026; the original is "
        "reversed by the credit memo of 6 April 2026.",
        "Rounding: entries a month to the nearest 50, entries per post to the nearest 5, hours to the hour and "
        "charges to the dollar; totals are summed before rounding.",
        "Prepared by Brokerage Workforce Planning for the establishment review of 22 October 2026.",
    ]
    for i, t in enumerate(notes, 1):
        c = nt.cell(row=i, column=1, value=t)
        c.font = bold if i == 1 else norm
        c.alignment = Alignment(wrap_text=True, vertical="top")
    nt.column_dimensions["A"].width = 120
    for sh in wb.worksheets:
        sh.page_setup.orientation = "landscape"
        sh.page_setup.paperSize = sh.PAPERSIZE_LETTER
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
        sh.sheet_properties.pageSetUpPr.fitToPage = True
        sh.print_options.horizontalCentered = True
        sh.page_margins.left = sh.page_margins.right = 0.5
    ws.print_title_rows = "4:4"
    wb.save(path)


# ---------------------------------------------------------------------------------------------------- the deck
def overtime_chart_png(G):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import rcParams
    rcParams["svg.hashsalt"] = "task145"
    rcParams["font.family"] = "DejaVu Sans"
    fig, ax = plt.subplots(figsize=(10.0, 4.6), dpi=160)
    fig.patch.set_facecolor("#fcfcfb")
    ax.set_facecolor("#fcfcfb")
    n = len(K.GATEWAYS)
    width = 0.15
    xs = list(range(4))
    for i, g in enumerate(K.GATEWAYS):
        vals = [G["A"][(g, q)] for q in (1, 2, 3, 4)]
        pos = [x + (i - (n - 1) / 2) * (width + 0.012) for x in xs]
        bars = ax.bar(pos, vals, width=width, color=SERIES[g], label=K.DESK[g], edgecolor="#fcfcfb", linewidth=1.0,
                      zorder=3)
        for b, v in zip(bars, vals):
            ax.text(b.get_x() + b.get_width() / 2, b.get_height() + 3, "{:,}".format(K.round_to(v, 1)), ha="center",
                    va="bottom", fontsize=7.4, color="#52514e")
    ax.set_xticks(xs)
    ax.set_xticklabels([FQ_LABEL[q] for q in (1, 2, 3, 4)], fontsize=9.5, color="#0b0b0b")
    ax.set_ylabel("Overtime hours worked", fontsize=9.5, color="#52514e")
    ax.tick_params(axis="y", labelsize=8.5, colors="#52514e")
    ax.yaxis.grid(True, color="#e3e2de", linewidth=0.8, zorder=0)
    for sp in ("top", "right", "left"):
        ax.spines[sp].set_visible(False)
    ax.spines["bottom"].set_color("#a3aab3")
    top = max(G["A"].values())
    ax.set_ylim(0, top * 1.18)
    ax.legend(ncol=5, frameon=False, fontsize=8.8, loc="upper left", bbox_to_anchor=(0.0, 1.10))
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor=fig.get_facecolor(), metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def chart_title(G):
    """The overtime slide's title: the finding the hours bear out (asserted in checks.py)."""
    peaks = {g: max((1, 2, 3, 4), key=lambda q: G["A"][(g, q)]) for g in K.GATEWAYS}
    tops = {q: max(K.GATEWAYS, key=lambda g: G["A"][(g, q)]) for q in (1, 2, 3, 4)}
    if all(p == 3 for p in peaks.values()):
        return "Overtime peaked at every desk in October to December 2025"
    if len(set(tops.values())) == 1:
        return "%s worked the most overtime of the five desks in every quarter of 2025-26" % K.DESK[tops[1]]
    return "Overtime hours worked at each desk, fiscal 2025-26"


def write_deck(path, G):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    blank = prs.slide_layouts[6]
    green = RGBColor(0x1F, 0x51, 0x30)
    grey = RGBColor(0x5A, 0x62, 0x70)
    ink = RGBColor(0x1D, 0x22, 0x29)

    def title(slide, text):
        tb = slide.shapes.add_textbox(Emu(457200), Emu(300000), Emu(11277600), Emu(820000))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = green
        p.font.name = "Calibri"

    def foot(slide, text, n):
        tb = slide.shapes.add_textbox(Emu(457200), Emu(6330000), Emu(10800000), Emu(380000))
        p = tb.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(9)
        p.font.color.rgb = grey
        p.font.name = "Calibri"
        nb = slide.shapes.add_textbox(Emu(11300000), Emu(6330000), Emu(500000), Emu(380000))
        q = nb.text_frame.paragraphs[0]
        q.text = str(n)
        q.font.size = Pt(9)
        q.font.color.rgb = grey

    s1 = prs.slides.add_slide(blank)
    title(s1, "Two of the six new broker posts go to Dorval and one to Calgary; three are held to the 2028 review")
    rows, cols = len(K.GATEWAYS) + 1, 5
    tbl = s1.shapes.add_table(rows, cols, Emu(457200), Emu(1450000), Emu(11277600), Emu(3300000)).table
    head = ["Desk", "Broker posts at 1 January 2027", "New posts placed", "Forecast 2027 formal entries, average "
            "month", "Formal entries a month per broker post after the review"]
    widths = [2400000, 2150000, 1900000, 2400000, 2427600]
    from pptx.enum.text import PP_ALIGN
    white = RGBColor(0xFF, 0xFF, 0xFF)
    pale = RGBColor(0xE3, 0xEB, 0xE5)
    for j, (h, w) in enumerate(zip(head, widths)):
        tbl.columns[j].width = Emu(w)
        c = tbl.cell(0, j)
        c.text = h
        c.fill.solid()
        c.fill.fore_color.rgb = green
        for p in c.text_frame.paragraphs:
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.name = "Calibri"
            p.font.color.rgb = white
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
    for i, g in enumerate(K.GATEWAYS, 1):
        vals = [K.DESK[g], str(K.ESTABLISHMENT[g]), str(G["placed"][g]),
                "{:,}".format(K.round_to(G["E"][g], 50)), str(K.round_to(G["final"][g], 5))]
        for j, v in enumerate(vals):
            c = tbl.cell(i, j)
            c.text = v
            c.fill.solid()
            c.fill.fore_color.rgb = pale if G["placed"][g] else white
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(14)
                p.font.name = "Calibri"
                p.font.color.rgb = ink
                p.font.bold = bool(G["placed"][g]) and j in (0, 2)
                p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
    nb = s1.shapes.add_textbox(Emu(457200), Emu(4950000), Emu(11277600), Emu(1100000))
    tf = nb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = ("Posts placed one at a time to the highest load while it exceeds the standard 560 (BRK-STD-03 4.5): "
              "Dorval, Calgary, Dorval. After three, the highest load is %s at %d, so three posts are held." % (
                  K.DESK[max(G["final"], key=G["final"].get)], K.round_to(max(G["final"].values()), 5)))
    p.font.size = Pt(13)
    p.font.color.rgb = ink
    p.font.name = "Calibri"
    p2 = tf.add_paragraph()
    p2.text = ("Forecast on the 2027 network plan: each gateway's shipments routed at the plan's exchange rates "
               "(USD 1.34, EUR 1.68), the line at CAD 3,300.")
    p2.font.size = Pt(12)
    p2.font.color.rgb = grey
    p2.font.name = "Calibri"
    foot(s1, "Brokerage Workforce Planning, establishment review 2027. Sources: BRK-STD-03 rev. 4; network plan 2027; "
             "high-value check extracts October 2025 to September 2026.", 1)

    s2 = prs.slides.add_slide(blank)
    title(s2, chart_title(G))
    png = overtime_chart_png(G)
    s2.shapes.add_picture(io.BytesIO(png), Emu(646000), Emu(1200000), width=Emu(10900000))
    foot(s2, "Overtime at time and a half and double time, and call-backs at the time worked, on the desk each broker "
             "was assigned to that day. Source: time and attendance exports, HR assignment changes.", 2)
    prs.save(path)
    return png


def main(target, golden_dir, figures_path, W=None, staff=None, lines=None, stmts=None):
    os.makedirs(golden_dir, exist_ok=True)
    G = golden(W, staff, lines, stmts)
    write_book(os.path.join(golden_dir, GOLD_BOOK), G)
    write_deck(os.path.join(golden_dir, GOLD_DECK), G)
    with open(figures_path, "w", encoding="utf-8") as fh:
        json.dump(figures(G), fh, indent=2, ensure_ascii=False)
        fh.write("\n")
    return G
