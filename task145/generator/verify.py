"""task145 independent verifier. Reads only the shipped bundle (target/) and recomputes, on its own code path, the call
(each segment's pass-through from the repeat consignments of the two years on file, the forecast at the 2027 plan's
rates with it, the loads and the placement), the default path's placement and the three rivals' (every shipment
re-routed at the plan's rates, one share for the network, one share per currency), and both asks; then compares them
with generator/figures.json and with the two golden files. It imports nothing from the generator: every
constant (the line, the standard load, the posts funded, the establishment, the plan's rates and volumes) is read from
the documents that file it.

    python3 task145/generator/verify.py [target] [golden]"""
from __future__ import annotations

import csv
import datetime as dt
import json
import math
import os
import re
import sys
from collections import defaultdict

import openpyxl

GW = ["YYZ", "YUL", "YVR", "YYC", "YWG"]
BASE = (dt.date(2025, 10, 1), dt.date(2026, 9, 30))
FY = (dt.date(2025, 4, 1), dt.date(2026, 3, 31))
HV = ["highvalue_check_2024Q4.csv", "highvalue_check_2025Q1.csv", "highvalue_check_2025Q2.csv",
      "highvalue_check_2025Q3.csv", "highvalue_check_2025Q4.csv", "highvalue_check_2026Q1.csv",
      "highvalue_check_2026Q2.csv", "highvalue_check_2026Q3.csv"]
PRIOR = (dt.date(2024, 10, 1), dt.date(2025, 9, 30))
MAIN = HV + ["gateway_activity_2024-10_2026-09.csv", "network_plan_2027.xlsx",
             "brk_std_03_brokerage_staffing_standard.pdf", "clr_ops_11_clearance_routing.docx"]
ASK_A = ["ta_legacy_export_2025-04-01_2025-10-31.csv", "ta_export_2025-10-27_2026-09-30.csv",
         "hr_brokerage_staff_2026-10-06.xlsx", "hr_assignment_changes_2025-04_2026-09.csv"]
ASK_B = ["ap_invoice_register_brandmoor_2025-04_2026-09.xlsx", "brandmoor_portal_detail_2025-04_2026-03.json",
         "finance_cost_centres_2026.csv"]
DESKS = ["Mississauga", "Dorval", "Richmond", "Calgary", "Winnipeg"]


def rnd(x, step):
    """Half away from zero, to a multiple of step."""
    q = x / step
    return int(math.floor(q + 0.5)) * step if q >= 0 else -int(math.floor(-q + 0.5)) * step


def cents(x):
    return math.floor(x * 100 + 0.5) / 100 if x >= 0 else -math.floor(-x * 100 + 0.5) / 100


def fq(d):
    return {4: 1, 5: 1, 6: 1, 7: 2, 8: 2, 9: 2, 10: 3, 11: 3, 12: 3, 1: 4, 2: 4, 3: 4}[d.month]


def pdf_text(path):
    from pypdf import PdfReader
    return re.sub(r"\s+", " ", " ".join(pg.extract_text() for pg in PdfReader(path).pages))


def docx_text(path):
    """A Word document's text, paragraph by paragraph, from its document part."""
    import html
    import zipfile
    with open(path, "rb") as fh, zipfile.ZipFile(fh) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    paras = []
    for p in re.findall(r"<w:p[ >].*?</w:p>", xml, re.S):
        p = re.sub(r"<w:tab/>", "<w:t> </w:t>", p)
        paras.append("".join(re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", p)))
    return html.unescape(re.sub(r"\s+", " ", " ".join(paras)))


# ---------------------------------------------------------------------------------------------------- the documents
def standard(target):
    """BRK-STD-03: the standard load, the posts funded, the establishment (appendix A) and the desks' names."""
    t = pdf_text(os.path.join(target, "brk_std_03_brokerage_staffing_standard.pdf"))
    bar = float(re.search(r"standard load of (\d+) formal entries a month", t).group(1))
    posts = int(re.search(r"New broker posts funded for 2027: (\d+)", t).group(1))
    app = t[t.index("Appendix A"):]
    est, desk = {}, {}
    for m in re.finditer(r"(Mississauga|Dorval|Richmond|Calgary|Winnipeg) (YYZ|YUL|YVR|YYC|YWG) (\d+)", app):
        est[m.group(2)] = int(m.group(3))
        desk[m.group(2)] = m.group(1)
    assert sorted(est) == sorted(GW), est
    return dict(bar=bar, posts=posts, est=est, desk=desk)


def line_cad(target):
    t = docx_text(os.path.join(target, "clr_ops_11_clearance_routing.docx"))
    return float(re.search(r"its value for duty exceeds CAD ([\d,]+)", t).group(1).replace(",", ""))


def plan(target):
    """The 2027 network plan: Treasury's planning rates (assumption 7) and the inbound volumes by gateway and month,
    thousands of shipments."""
    wb = openpyxl.load_workbook(os.path.join(target, "network_plan_2027.xlsx"), read_only=True)
    rates = {"CAD": 1.0}
    for r in wb["Assumptions"].iter_rows(values_only=True):
        if r[1] == "Exchange rates":
            for c, v in re.findall(r"\b(USD|EUR|GBP|CNY|JPY) (\d+\.\d+)", r[2]):
                rates[c] = float(v)
    vols = {}
    for r in wb["Inbound volumes"].iter_rows(values_only=True):
        m = re.search(r"\((YYZ|YUL|YVR|YYC|YWG)\)$", str(r[0]))
        if m:
            vols[m.group(1)] = [float(x) for x in r[1:13]]
    assert len(rates) == 6 and sorted(vols) == sorted(GW), (rates, vols)
    return rates, vols


def gateway_report(target):
    """Base-year inbound international shipments and formal entries released, by gateway (current report)."""
    out = {g: [0, 0] for g in GW}
    with open(os.path.join(target, "gateway_activity_2024-10_2026-09.csv"), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            y, m = map(int, r["period"].split("-"))
            if BASE[0] <= dt.date(y, m, 1) <= BASE[1]:
                out[r["gateway_code"]][0] += int(r["inbound_international_shipments"])
                out[r["gateway_code"]][1] += int(r["formal_entries_released"])
    return out


def register(target, keep="last"):
    """The eight high-value check extracts in their two layouts, one record per shipment: an amended shipment's row
    in the later extract replaces the earlier (keep='both' keeps every row)."""
    recs = {}
    allrows = []
    for f in HV:
        with open(os.path.join(target, f), encoding="utf-8") as fh:
            for r in csv.DictReader(fh):
                if "WAYBILL" in r:
                    rec = dict(wb=r["WAYBILL"], gw=r["GTWY"], local=dt.datetime.strptime(r["REL_DT"], "%d/%m/%Y").date(),
                               utc=r["REL_TM_UTC"], imp=r["IMPORTER_ACCT"], hs=r["HS_CODE"], cur=r["INV_CURR"],
                               val=float(r["VFD_INV"]),
                               rate=float(r["CUST_RATE"]), cad=float(r["VFD_CAD"]), route=r["ROUTE"],
                               pgm=r["OGD_PGM"].strip(), amended=r["AMEND"] == "Y", file=f)
                else:
                    rec = dict(wb=r["waybill"], gw=r["gateway"], local=dt.date.fromisoformat(r["release_local_date"]),
                               utc=r["release_utc"], imp=r["importer_account"], hs=r["hs_code"],
                               cur=r["invoice_currency"], val=float(r["value_for_duty"]),
                               rate=float(r["customs_exchange_rate"]), cad=float(r["value_for_duty_cad"]),
                               route=r["route"], pgm=r["ogd_program"].strip(), amended=r["amended"] == "Y", file=f)
                rec["ogd"] = rec["pgm"] not in ("", "NONE")
                allrows.append(rec)
                recs[rec["wb"]] = rec
    return allrows if keep == "both" else list(recs.values())


def formal_count(reg, rates, line):
    """Each shipment routed as CLR-OPS-11 section 3 routes it, at the given rate for its invoice currency."""
    n = {g: 0 for g in GW}
    for s in reg:
        if s["ogd"] or cents(s["val"] * rates[s["cur"]]) > line:
            n[s["gw"]] += 1
    return n


def place(E, est, bar, posts):
    """BRK-STD-03 4.5: one post at a time to the highest load, the posts already placed counted, while it exceeds the
    standard load; the rest held."""
    placed = {g: 0 for g in GW}
    order = []
    for _ in range(posts):
        load = {g: E[g] / (est[g] + placed[g]) for g in GW}
        g = max(load, key=load.get)
        if load[g] <= bar:
            break
        placed[g] += 1
        order.append((g, load[g]))
    final = {g: E[g] / (est[g] + placed[g]) for g in GW}
    return placed, posts - sum(placed.values()), final, order


def in_base(s):
    return BASE[0] <= s["local"] <= BASE[1]


def repeats(reg_all, seg=lambda s: (s["gw"], s["cur"])):
    """How the goods' values behave between the two years on file, by segment (gateway and invoice currency unless
    seg says otherwise): the base year's foreign-currency shipments outside OGD programs that repeat a consignment of
    the year before (the same gateway, importer, tariff item and currency) at the same invoice value, and those that
    repeat one at the same CAD value. A consignment that keeps its invoice value passes the exchange rate through to its
    CAD value; one that keeps its CAD value does not."""
    fx = [s for s in reg_all if not s["ogd"] and s["cur"] != "CAD"]
    prior = [s for s in fx if PRIOR[0] <= s["local"] <= PRIOR[1]]
    k_inv = lambda s: (s["gw"], s["imp"], s["hs"], s["cur"], s["val"])
    k_cad = lambda s: (s["gw"], s["imp"], s["hs"], s["cur"], s["cad"])
    p_inv, p_cad = {k_inv(s) for s in prior}, {k_cad(s) for s in prior}
    out = defaultdict(lambda: [0, 0])
    for s in fx:
        if not in_base(s):
            continue
        if k_inv(s) in p_inv:
            out[seg(s)][0] += 1
        elif k_cad(s) in p_cad:
            out[seg(s)][1] += 1
    return {k: (a, b, a / (a + b)) for k, (a, b) in out.items()}


def with_passthrough(reg, rates, line, share, seg=lambda s: (s["gw"], s["cur"])):
    """Each gateway's base-year shipments routed for 2027: OGD goods formal; goods invoiced in CAD on their value; in
    every other segment its pass-through share re-routed at the plan's rate on the invoice value, the rest at the CAD
    value (and so the route) the base year gave them."""
    n = {g: 0.0 for g in GW}
    for s in reg:
        if s["ogd"] or s["cur"] == "CAD":
            n[s["gw"]] += s["route"] == "FORMAL"
            continue
        p = share(seg(s))
        n[s["gw"]] += p * (cents(s["val"] * rates[s["cur"]]) > line) + (1 - p) * (s["route"] == "FORMAL")
    return n


def main_call(target):
    st = standard(target)
    line = line_cad(target)
    rates, vols = plan(target)
    rep = gateway_report(target)
    reg_all = register(target)
    reg = [s for s in reg_all if in_base(s)]
    ratio = {g: sum(vols[g]) * 1000.0 / rep[g][0] for g in GW}
    scale = lambda n: {g: n[g] * ratio[g] / 12.0 for g in GW}
    pt = repeats(reg_all)
    n = with_passthrough(reg, rates, line, lambda k: pt[k][2])
    E = scale(n)
    placed, held, final, order = place(E, st["est"], st["bar"], st["posts"])
    # (a) the default: the base year's formal entries per shipment on the plan's volumes
    Ed = {g: rep[g][1] / rep[g][0] * sum(vols[g]) * 1000.0 / 12.0 for g in GW}
    pd_, hd, fd, od = place(Ed, st["est"], st["bar"], st["posts"])
    # (b) every foreign-currency shipment re-routed at the plan's rate (pass-through one everywhere)
    nb = formal_count(reg, rates, line)
    Eb = scale(nb)
    pb, hb, fb, ob = place(Eb, st["est"], st["bar"], st["posts"])
    # (c) one pass-through share for the network; (d) one for each invoice currency
    ptc = repeats(reg_all, seg=lambda s: "all")
    Ec = scale(with_passthrough(reg, rates, line, lambda k: ptc["all"][2]))
    pc, hc, fc, oc = place(Ec, st["est"], st["bar"], st["posts"])
    ptd = repeats(reg_all, seg=lambda s: s["cur"])
    Edc = scale(with_passthrough(reg, rates, line, lambda k: ptd[k][2], seg=lambda s: s["cur"]))
    pdc, hdc, fdc, odc = place(Edc, st["est"], st["bar"], st["posts"])
    return dict(st=st, line=line, rates=rates, vols=vols, rep=rep, reg=reg, reg_all=reg_all, ratio=ratio, n=n, E=E,
                placed=placed, held=held, final=final, order=order, Ed=Ed, placed_d=pd_, held_d=hd, final_d=fd,
                nb=nb, Eb=Eb, placed_b=pb, held_b=hb, final_b=fb, pt=pt, ptc=ptc, Ec=Ec, placed_c=pc, held_c=hc,
                final_c=fc, ptd=ptd, Edc=Edc, placed_dc=pdc, held_dc=hdc, final_dc=fdc)


# ---------------------------------------------------------------------------------------------------- Ask A
def staff(target):
    wb = openpyxl.load_workbook(os.path.join(target, "hr_brokerage_staff_2026-10-06.xlsx"), read_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    ix = {h: i for i, h in enumerate(rows[0])}
    out = {}
    for r in rows[1:]:
        if not r[ix["Employee number"]]:
            continue
        out[r[ix["Employee number"]]] = dict(legacy=r[ix["Legacy ID"]], desk=r[ix["Desk"]],
                                             position=r[ix["Position"]])
    return out


def desk_timeline(target, emp):
    """Each employee's desk by date from the HR assignment changes: the hire desk or the desk before the first
    transfer, then each transfer from its effective date."""
    ch = defaultdict(list)
    with open(os.path.join(target, "hr_assignment_changes_2025-04_2026-09.csv"), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            if r["change_type"] in ("HIRE", "TRANSFER"):
                ch[r["employee_number"]].append((dt.date.fromisoformat(r["effective_date"]), r["change_type"],
                                                 r["from_desk"], r["to_desk"]))
    tl = {}
    for e, info in emp.items():
        evs = sorted(ch.get(e, []))
        start = (evs[0][2] if evs[0][1] == "TRANSFER" else evs[0][3]) if evs else info["desk"]
        tl[e] = (start, [(d, to) for d, _, _, to in evs])
    return tl


def desk_on(tl, e, day):
    desk, evs = tl[e]
    for d, to in evs:
        if d <= day:
            desk = to
    return desk


def time_lines(target, emp):
    """Overtime and call-back lines from both exports: the legacy export to 31 October 2025, the current system's own
    lines after it (its MIGRATED lines are the legacy system's timecards of 27 to 31 October, copied)."""
    by_legacy = {v["legacy"]: k for k, v in emp.items()}
    out = []
    with open(os.path.join(target, "ta_legacy_export_2025-04-01_2025-10-31.csv"), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            code = {"OT15": "OT", "OT20": "OT", "CALLB": "CB"}.get(r["PAY_CODE"])
            if not code:
                continue
            day = dt.datetime.strptime(r["WORK_DATE"], "%m/%d/%Y").date()
            a = dt.datetime.combine(day, dt.datetime.strptime(r["IN_TIME"], "%H:%M").time())
            b = dt.datetime.combine(day, dt.datetime.strptime(r["OUT_TIME"], "%H:%M").time())
            if b <= a:
                b += dt.timedelta(days=1)
            out.append(dict(emp=by_legacy[r["EMPLID"]], day=day, code=code, hours=float(r["HOURS"]),
                            worked=(b - a).total_seconds() / 3600.0, src="LEGACY"))
    with open(os.path.join(target, "ta_export_2025-10-27_2026-09-30.csv"), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            code = {"OT1.5": "OT", "OT2.0": "OT", "CB": "CB"}.get(r["earning_code"])
            if not code:
                continue
            a = dt.datetime.fromisoformat(r["punch_in"])
            b = dt.datetime.fromisoformat(r["punch_out"])
            out.append(dict(emp=r["employee_number"], day=dt.date.fromisoformat(r["work_date"]), code=code,
                            hours=float(r["quantity"]), worked=(b - a).total_seconds() / 3600.0, src=r["source"]))
    return out


def ask_a(target):
    """Overtime hours worked, by desk name and fiscal quarter: overtime lines at their hours, call-backs at the time
    between their punches, each line on the desk the employee held on the work date."""
    emp = staff(target)
    tl = desk_timeline(target, emp)
    out = {(n, q): 0.0 for n in DESKS for q in (1, 2, 3, 4)}
    for ln in time_lines(target, emp):
        if ln["src"] == "MIGRATED" or not (FY[0] <= ln["day"] <= FY[1]):
            continue
        h = ln["worked"] if ln["code"] == "CB" else ln["hours"]
        out[(desk_on(tl, ln["emp"], ln["day"]), fq(ln["day"]))] += h
    return out


# ---------------------------------------------------------------------------------------------------- Ask B
def ap_register(target):
    wb = openpyxl.load_workbook(os.path.join(target, "ap_invoice_register_brandmoor_2025-04_2026-09.xlsx"),
                                read_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    ix = {h: i for i, h in enumerate(rows[0])}
    out = []
    for r in rows[1:]:
        if not r[ix["Document number"]]:
            continue
        d = r[ix["Document date"]]
        p = r[ix["Posting date"]]
        out.append(dict(type=r[ix["Document type"]], number=r[ix["Document number"]],
                        date=d.date() if hasattr(d, "date") else d, post=p.date() if hasattr(p, "date") else p,
                        cc=str(r[ix["Cost centre"]]), net=float(r[ix["Net CAD"]]), ref=r[ix["Reference"]]))
    return out


def cost_centres(target):
    with open(os.path.join(target, "finance_cost_centres_2026.csv"), encoding="utf-8") as fh:
        return {r["cost_centre"]: r["desk"] for r in csv.DictReader(fh) if r["desk"]}


def portal_lines(target):
    with open(os.path.join(target, "brandmoor_portal_detail_2025-04_2026-03.json"), encoding="utf-8") as fh:
        return json.load(fh)["lines"]


def portal_months(target):
    """Each statement's transmission months, from Brandmoor's portal detail."""
    out = defaultdict(set)
    for r in portal_lines(target):
        t = dt.datetime.fromisoformat(r["transmitted_et"])
        out[r["statement_number"]].add((t.year, t.month))
    return out


def prev_month(d):
    return (d.year - 1, 12) if d.month == 1 else (d.year, d.month - 1)


def ask_b(target):
    """Each statement's net before tax on the month of the transmissions it bills (the month before its date, as the
    portal shows for every statement it holds), a statement reversed by a credit memo replaced by its re-issue."""
    ap = ap_register(target)
    cc = cost_centres(target)
    pm = portal_months(target)
    reversed_ = {re.search(r"Reverses (\S+)", r["ref"]).group(1) for r in ap if r["type"] == "Credit memo"}
    out = {(n, q): 0.0 for n in DESKS for q in (1, 2, 3, 4)}
    for r in ap:
        if r["type"] != "Invoice" or r["number"] in reversed_:
            continue
        ym = prev_month(r["date"])
        if r["number"] in pm:
            assert pm[r["number"]] == {ym}, (r["number"], pm[r["number"]], ym)
        d = dt.date(ym[0], ym[1], 1)
        if FY[0] <= d <= FY[1]:
            out[(cc[r["cc"]], fq(d))] += r["net"]
    return out


# ---------------------------------------------------------------------------------------------------- comparison
def fmt_int(x, step):
    return "{:,}".format(rnd(x, step))


def golden_values(golden):
    """The figures as the golden workbook and deck carry them."""
    wb = openpyxl.load_workbook(os.path.join(golden, "broker_posts_2027.xlsx"), read_only=True)
    pl = list(wb["Placement"].iter_rows(values_only=True))
    rows = {r[1]: r for r in pl if r and r[1] in GW}
    held = [r for r in pl if r and r[0] == "New posts held to the 2028 review"][0][3]
    ot = {r[0]: r[1:5] for r in wb["Overtime FY2025-26"].iter_rows(values_only=True) if r and r[0] in
          ("Mississauga", "Dorval", "Richmond", "Calgary", "Winnipeg")}
    bm = {r[0]: r[1:5] for r in wb["Brandmoor FY2025-26"].iter_rows(values_only=True) if r and r[0] in
          ("Mississauga", "Dorval", "Richmond", "Calgary", "Winnipeg")}
    basis_seg, basis_gw = {}, {}
    for r in wb["Forecast basis"].iter_rows(values_only=True):
        if not r or r[0] not in DESKS:
            continue
        if isinstance(r[1], str):
            basis_seg[(r[0], r[1])] = tuple(r[2:5])
        else:
            basis_gw[r[0]] = tuple(r[1:6])
    from pptx import Presentation
    prs = Presentation(os.path.join(golden, "establishment_review_2027.pptx"))
    tbl = [sh.table for sh in prs.slides[0].shapes if sh.has_table][0]
    deck = {tbl.cell(i, 0).text: [tbl.cell(i, j).text for j in range(1, 5)] for i in range(1, len(tbl.rows))}
    return rows, held, ot, bm, deck, basis_seg, basis_gw


def main(target=None, golden=None):
    here = os.path.dirname(os.path.abspath(__file__))
    task = os.path.dirname(here)
    target = target or os.path.join(task, "target")
    golden = golden or os.path.join(task, "golden")
    fig = json.load(open(os.path.join(here, "figures.json"), encoding="utf-8"))
    c = main_call(target)
    A = ask_a(target)
    B = ask_b(target)
    d = c["st"]["desk"]
    bad = []

    def eq(label, got, want):
        if got != want:
            bad.append("%s: bundle %r, figures %r" % (label, got, want))

    eq("posts placed", {d[g]: str(c["placed"][g]) for g in GW}, fig["posts_placed"])
    eq("held", str(c["held"]), fig["call"]["held"])
    eq("forecast", {d[g]: fmt_int(c["E"][g], 50) for g in GW}, fig["forecast_monthly_formal_entries"])
    eq("entries per post", {d[g]: str(rnd(c["final"][g], 5)) for g in GW}, fig["entries_per_post_after"])
    eq("establishment", {d[g]: str(c["st"]["est"][g]) for g in GW}, fig["establishment"])
    eq("default placement", {d[g]: str(c["placed_d"][g]) for g in GW}, fig["default_path"]["posts_placed"])
    eq("default held", str(c["held_d"]), fig["default_path"]["held"])
    pt = c["pt"]
    eq("pass-through", {"%s %s" % (d[g], cu): {"repeats_on_invoice_value": "{:,}".format(a),
                                                 "repeats_on_cad_value": "{:,}".format(b), "share": "%.3f" % p}
                         for (g, cu), (a, b, p) in sorted(pt.items(), key=lambda kv: (GW.index(kv[0][0]), kv[0][1]))},
       fig["pass_through"])
    eq("pooled share", "%.3f" % c["ptc"]["all"][2], fig["rivals"]["pooled_share"]["share"])
    eq("US dollar share", "%.3f" % c["ptd"]["USD"][2], fig["rivals"]["by_currency"]["usd_share"])
    for key, p_, h_, E_ in (("full_repricing", c["placed_b"], c["held_b"], c["Eb"]),
                            ("pooled_share", c["placed_c"], c["held_c"], c["Ec"]),
                            ("by_currency", c["placed_dc"], c["held_dc"], c["Edc"])):
        eq("rival %s placement" % key, {d[g]: str(p_[g]) for g in GW}, fig["rivals"][key]["posts_placed"])
        eq("rival %s held" % key, str(h_), fig["rivals"][key]["held"])
        eq("rival %s forecast" % key, {d[g]: fmt_int(E_[g], 50) for g in GW},
           fig["rivals"][key]["forecast_monthly_formal_entries"])
    eq("formal entries with the pass-through", {d[g]: "{:,}".format(rnd(c["n"][g], 1)) for g in GW},
       fig["base_year_formal_with_pass_through"])
    eq("formal entries at the plan's rates", {d[g]: "{:,}".format(c["nb"][g]) for g in GW},
       fig["base_year_formal_at_plan_rates"])
    eq("ask A", {d[g]: {"FQ%d" % q: fmt_int(A[(d[g], q)], 1) for q in (1, 2, 3, 4)} for g in GW},
       fig["ask_a_overtime_hours"])
    eq("ask B", {d[g]: {"FQ%d" % q: fmt_int(B[(d[g], q)], 1) for q in (1, 2, 3, 4)} for g in GW},
       fig["ask_b_brandmoor_cad"])
    rows, held, ot, bm, deck, basis_seg, basis_gw = golden_values(golden)
    eq("workbook pass-through", basis_seg, {(d[g], cu): (a, b, round(p, 3)) for (g, cu), (a, b, p) in c["pt"].items()})
    n_base = {g: sum(1 for s in c["reg"] if s["gw"] == g and s["route"] == "FORMAL") for g in GW}
    eq("workbook forecast basis", basis_gw, {d[g]: (n_base[g], c["nb"][g], rnd(c["n"][g], 1), round(c["ratio"][g], 4),
                                                    rnd(c["E"][g], 50)) for g in GW})
    eq("formal entries as routed", {d[g]: "{:,}".format(n_base[g]) for g in GW}, fig["base_year_formal_as_routed"])
    eq("plan over base", {d[g]: "%.4f" % c["ratio"][g] for g in GW}, fig["plan_over_base"])
    for g in GW:
        r = rows[g]
        eq("workbook %s" % g, (r[3], r[5], r[6]), (c["placed"][g], rnd(c["E"][g], 50), rnd(c["final"][g], 5)))
        eq("workbook overtime %s" % g, tuple(ot[d[g]]), tuple(rnd(A[(d[g], q)], 1) for q in (1, 2, 3, 4)))
        eq("workbook Brandmoor %s" % g, tuple(bm[d[g]]), tuple(rnd(B[(d[g], q)], 1) for q in (1, 2, 3, 4)))
        eq("deck %s" % g, deck[d[g]], [str(c["st"]["est"][g]), str(c["placed"][g]), fmt_int(c["E"][g], 50),
                                        str(rnd(c["final"][g], 5))])
    eq("workbook held", held, c["held"])
    if bad:
        print("verify.py: %d mismatches" % len(bad))
        for b in bad:
            print("  " + b)
        raise SystemExit(1)
    print("verify.py: from the bundle alone, %s; held %d; forecast %s; default path %s, held %d; full repricing held "
          "%d, one network share held %d, one share per currency held %d; %d overtime and %d Brandmoor figures; all "
          "equal figures.json, the workbook and the deck" % (
              ", ".join("%s %d" % (d[g], c["placed"][g]) for g in GW), c["held"],
              ", ".join("%s %s" % (d[g], fmt_int(c["E"][g], 50)) for g in GW),
              ", ".join("%s %d" % (d[g], c["placed_d"][g]) for g in GW), c["held_d"], c["held_b"], c["held_c"],
              c["held_dc"], len(A), len(B)))
    return c, A, B


if __name__ == "__main__":
    main(*(sys.argv[1:3]))
