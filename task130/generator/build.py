#!/usr/bin/env python3
"""task130 generator: builds the evidence pack into <out>/target and <out>/metadata.json, asserting
everything the design note's assertion plan lists.

    python3 build.py --out /path/to/task130 [--record /path/to/record.json]
"""
import argparse
import datetime as dt
import json
import os
import random
import shutil
import sys

sys.dont_write_bytecode = True
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import analysis as A  # noqa: E402
import core  # noqa: E402
import distractors  # noqa: E402
import docs  # noqa: E402
import ledger  # noqa: E402
import padro  # noqa: E402
import returns  # noqa: E402
import tables  # noqa: E402
import writers as WR  # noqa: E402
from common import SEED, r1  # noqa: E402
from plan import SECTIONS  # noqa: E402
from world import GROUP_CODE, GROUP_WORD  # noqa: E402

DT = dt.datetime
F = {
    "returns": "declaracions_grans_tenidors_2024T3_2026T2.csv",
    "guide": "guia_presentacio_declaracio_trimestral_2024-01.pdf",
    "notice": "avis_portal_canvi_format_2026-03.eml",
    "register": "registre_grans_tenidors_20261001.xlsx",
    "codes": "codis_grup.csv",
    "cadastre": "cadastre_habitatges_extracte_20260928.csv",
    "deeds": "escriptures_inscrites_2024-07_2026-09.csv",
    "ledger": "EPPR_execucio_pressupost_2026_gen-set.csv",
    "bull2": "butlleti_parc_residencial_2026T2.pdf",
    "bull1": "butlleti_parc_residencial_2026T1.pdf",
    "annex26": "annex_designacio_2026.xlsx",
    "order": "ordre_designacio_seccions_grans_tenidors.pdf",
    "padro09": "padro_lliurament_2026-09.csv",
    "padro06": "padro_lliurament_2026-06.csv",
    "xnote": "nota_intercanvi_padro_rev2.txt",
    "note": "notes_annex_2026_CGallart.docx",
    "deposits": "fiances_lloguer_seccions_2025.csv",
    "visits": "inspeccions_habitatge_buit_2025.csv",
    "index": "index_expedient_designacio_2027.csv",
}
DISTRACTORS = [F["deposits"], F["visits"], F["bull1"]]
MTIME = {"returns": DT(2026, 10, 1, 8, 42), "guide": DT(2024, 1, 22, 12, 5), "notice": DT(2026, 3, 16, 9, 45),
         "register": DT(2026, 10, 1, 10, 17), "codes": DT(2026, 10, 1, 10, 18), "cadastre": DT(2026, 9, 28, 16, 3),
         "deeds": DT(2026, 10, 2, 11, 26), "ledger": DT(2026, 10, 5, 13, 51), "bull2": DT(2026, 9, 15, 10, 30),
         "bull1": DT(2026, 6, 16, 10, 12), "annex26": DT(2025, 11, 12, 17, 40), "order": DT(2025, 9, 3, 9, 15),
         "padro09": DT(2026, 9, 14, 8, 2), "padro06": DT(2026, 6, 12, 8, 4), "xnote": DT(2025, 3, 10, 13, 20),
         "note": DT(2026, 9, 24, 18, 36), "deposits": DT(2026, 3, 2, 9, 8), "visits": DT(2026, 1, 20, 15, 44),
         "index": DT(2026, 10, 6, 12, 10)}
WATCH = sorted(s["code"] for s in SECTIONS)


def read_xlsx_sheet(path, sheet):
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True)[sheet]
    rows = list(ws.iter_rows(values_only=True))
    head = rows[0]
    return [{h: ("" if v is None else str(v)) for h, v in zip(head, r)} for r in rows[1:]]


def bulletin_rows(ret, cad, quarter, cutoff, prev=None):
    c = A.counts(A.picture(ret, quarter, cutoff, cad, "R2"), WATCH)
    out = []
    for s in WATCH:
        sh = r1(A.share(c[s], cad.N[s]))
        dl = None
        if prev is not None:
            dl = sh - r1(A.share(prev[s], cad.N[s]))
        out.append((s, cad.N[s], c[s], sh, dl))
    return out, c


def write_data(W, T):
    p = lambda k: os.path.join(T, F[k])
    plan_mal = {}
    rrows = returns.render(W["lodgements"], W["holders"], W["stock"], W["nif"], plan_mal)
    W["plan_mal"] = plan_mal
    info = {"returns_rows": WR.write_csv(p("returns"), returns.HEADER, rrows)}
    on_returns = {x[0] for l in W["lodgements"] for x in l["rows"]}
    scope = tables.scope_refs(W, on_returns)
    W["scope"] = scope
    info["deeds_rows"] = WR.write_csv(p("deeds"), tables.DEED_HEADER, tables.deeds(W, scope))
    info["cad_rows"] = WR.write_csv(p("cadastre"), tables.CAD_HEADER, tables.cadastre(W, scope))
    tit, links = tables.register(W)
    docs.register_xlsx(p("register"), tit, links, "1 d'octubre de 2026")
    WR.write_csv(p("codes"), ["codi_grup", "nom_grup", "data_alta", "data_baixa"], tables.group_codes())
    info["ledger_rows"] = WR.write_csv(p("ledger"), ledger.HEADER, ledger.build(W))
    return info


def load(T):
    p = lambda k: os.path.join(T, F[k])
    L = {"ret": A.read_csv(p("returns")), "cad": A.Cadastre(A.read_csv(p("cadastre"))),
         "deeds": A.read_csv(p("deeds")), "led": A.read_csv(p("ledger")), "codes_rows": A.read_csv(p("codes")),
         "tit": read_xlsx_sheet(p("register"), "Titulars"), "links": read_xlsx_sheet(p("register"), "Vinculacions")}
    L["lh"] = {r["nif"] for r in L["tit"] if not r["data_baixa"]}
    L["names"] = {r["nif"]: r["nom"] for r in L["tit"]}
    L["codes"] = {r["codi_grup"]: (r["nom_grup"], r["data_alta"]) for r in L["codes_rows"]}
    return L


def compute(L):
    """Every rung, the 2025 close-out and the golden holdings, from the shipped tables."""
    ret, cad, deeds, led, lh = L["ret"], L["cad"], L["deeds"], L["led"], L["lh"]
    C = {}
    for m in ("R0", "R1", "R2"):
        C[m] = A.counts(A.picture(ret, "2026T2", "2026-08-31", cad, m), WATCH)
    own, sch = A.holdings(A.picture(ret, "2026T2", "2026-08-31", cad, "R2"))
    tk = A.ledger_takeovers(led, 120)
    rolls = {}
    for m in ("R3", "R4", "R4p", "R5"):
        rolls[m] = A.roll(own, sch, deeds, lh, "2026-06-30", "2026-09-30", "2026-09-30", "2027-01-01", tk, m)
        C[m] = A.section_counts(rolls[m], cad, WATCH)
    own25, sch25 = A.holdings(A.picture(ret, "2025T2", "2025-08-31", cad, "R2"))
    roll25 = A.roll(own25, sch25, deeds, lh, "2025-06-30", "2025-09-30", "2025-09-30", "2026-01-01")
    return {"C": C, "own": own, "sch": sch, "tk": tk, "rolls": rolls, "own25": own25, "sch25": sch25,
            "roll25": roll25, "c25": A.section_counts(roll25, cad, WATCH)}


def annex26_rows(L, R, W):
    cad = L["cad"]
    member = A.links_at(L["links"], "2026-01-01")
    old_names = {GROUP_CODE[g]: f"Grup {GROUP_WORD[g]}" for g in ("K1", "K2")}
    codes26 = {c: ((old_names.get(c) or v[0]), v[1]) for c, v in L["codes"].items()}
    big = A.largest(R["roll25"], cad, WATCH, member, codes26, L["names"])
    rng = random.Random(SEED * 71 + 1)
    rows = []
    for s in WATCH:
        n = R["c25"][s]
        sh = r1(A.share(n, cad.N[s]))
        rows.append((s, cad.N[s], n, sh, A.share(n, cad.N[s]) >= 25, big[s][0], big[s][1],
                     int(round(n * rng.uniform(0.045, 0.085)))))
    return rows, big


def write_docs(W, T, L, R):
    p = lambda k: os.path.join(T, F[k])
    ret, cad = L["ret"], L["cad"]
    q4, c4 = bulletin_rows(ret, cad, "2025T4", "2026-02-28")
    q1, c1 = bulletin_rows(ret, cad, "2026T1", "2026-05-31", c4)
    q2, c2 = bulletin_rows(ret, cad, "2026T2", "2026-08-31", c1)
    docs.bulletin_pdf(p("bull1"), "Primer trimestre de 2026", "31 de març de 2026", "31 de maig de 2026", q1,
                      "31/12/2025", "2026/1")
    docs.bulletin_pdf(p("bull2"), "Segon trimestre de 2026", "30 de juny de 2026", "31 d'agost de 2026", q2,
                      "31/03/2026", "2026/2")
    a_rows, big26 = annex26_rows(L, R, W)
    docs.annex_xlsx(p("annex26"), a_rows)
    docs.order_pdf(p("order"))
    docs.guide_pdf(p("guide"))
    WR.write_text(p("notice"), docs.notice_eml(), crlf=True)
    WR.write_text(p("xnote"), docs.EXCHANGE_NOTE)
    docs.note_docx(p("note"))
    return {"bull2": q2, "bull1": q1, "bull_q4": q4, "annex26": a_rows, "big26": big26}


def write_padro(W, T, R):
    touched = set(R["tk"]) | {r for a in W["book"].agreements if a["agreed"].isoformat() == "2026-12-10"
                              for r in a["refs"]}
    rows, status = padro.build(W["stock"], R["rolls"]["R5"], set(R["own"]), touched)
    WR.write_csv(os.path.join(T, F["padro06"]), padro.HEADER, rows["06"])
    WR.write_csv(os.path.join(T, F["padro09"]), padro.HEADER, rows["09"])
    return status, touched


INDEX_ROWS = [
    ("returns", "Declaracions trimestrals d'habitatges de grans tenidors, totes les versions, trimestres 2024T3 a 2026T2",
     "Portal de declaracions (Observatori)", "Declaracions presentades fins a l'extracció"),
    ("guide", "Guia de presentació de la declaració trimestral, versió 3", "Servei de Declaracions", ""),
    ("notice", "Avís del portal sobre el canvi de format (abril de 2026)", "Servei de Declaracions", ""),
    ("register", "Registre de Grans Tenidors: titulars, vinculacions a grups i diccionari", "Registre de Grans Tenidors",
     "Situació a la data d'extracció, amb les vinculacions anotades"),
    ("codes", "Taula de codis de grup", "Registre de Grans Tenidors", "Codis vigents i donats de baixa"),
    ("cadastre", "Extracte cadastral: unitats de les seccions de la llista de seguiment i habitatges que figuren en "
                 "alguna declaració", "Cadastre (extracció de l'Observatori)", "Unitats residencials i no residencials"),
    ("deeds", "Escriptures inscrites, juliol de 2024 a setembre de 2026", "Registres de la Propietat (extracció)",
     "Habitatges de les seccions de seguiment i habitatges que figuren en alguna declaració; inscrites fins al 30/09/2026"),
    ("ledger", "Execució del pressupost de l'Ens Públic de Patrimoni Residencial, gener a setembre de 2026",
     "Ens Públic de Patrimoni Residencial (L. Sacristán)", "Totes les fases i programes; comptabilitzat fins al 30/09/2026"),
    ("bull2", "Butlletí del Parc Residencial 2026/2", "Observatori del Parc Residencial", "Tinença a 30/06/2026"),
    ("bull1", "Butlletí del Parc Residencial 2026/1", "Observatori del Parc Residencial", "Tinença a 31/03/2026"),
    ("annex26", "Annex de l'ordre de designació 2026", "Observatori del Parc Residencial", "Situació a 01/01/2026"),
    ("order", "Ordre de designació, text consolidat", "Direcció General d'Habitatge", ""),
    ("padro09", "Lliurament del padró municipal, setembre de 2026", "Ajuntaments", "Seccions de seguiment"),
    ("padro06", "Lliurament del padró municipal, juny de 2026", "Ajuntaments", "Seccions de seguiment"),
    ("xnote", "Nota de protocol de l'intercanvi del padró, revisió 2", "Observatori del Parc Residencial", ""),
    ("note", "Notes de C. Gallart sobre l'annex de 2026", "Observatori del Parc Residencial", ""),
    ("deposits", "Registre de fiances: contractes de lloguer dipositats per secció, 2025", "Registre de fiances",
     "Any 2025"),
    ("visits", "Visites d'inspecció a habitatges de grans tenidors, 2025", "Inspecció d'Habitatge", "Any 2025"),
]


def write_index(T):
    rows = []
    for k, what, src, cov in INDEX_ROWS:
        rows.append([F[k], what, src, cov, MTIME[k].date().isoformat()])
    WR.write_csv(os.path.join(T, F["index"]), ["fitxer", "contingut", "origen", "cobertura", "data_fitxer"], rows)


def finish_containers(T):
    p = lambda k: os.path.join(T, F[k])
    for k, author in (("order", "Direccio General d'Habitatge"), ("guide", "Observatori Parc Residencial"),
                      ("bull1", "Observatori Parc Residencial"), ("bull2", "Observatori Parc Residencial")):
        WR.scrub_pdf(p(k), author, MTIME[k].date().isoformat())
    WR.normalize_ooxml(p("register"), "Registre de Grans Tenidors", MTIME["register"])
    WR.normalize_ooxml(p("annex26"), "Celestina Gallart", DT(2025, 11, 10, 11, 5), MTIME["annex26"])
    WR.normalize_ooxml(p("note"), "Celestina Gallart", DT(2026, 9, 24, 17, 52), MTIME["note"], app="Microsoft Office Word")
    for k, when in MTIME.items():
        WR.set_mtime(p(k), when)


def metadata(out, T, info):
    files = sorted(os.listdir(T))
    meta = {"task": "task130", "title": "2027 large-holder designation of census sections",
            "domain": "Demographic & Social Science", "subdomain": "housing",
            "objective": "Data Extraction & Conformation (ETL)", "as_of": "2026-09-30",
            "deliverables": ["designation_brief_2027.pdf", "designation_annex_2027.csv", "section_bridge_2027.png"],
            "input_files": [{"path": f, "format": os.path.splitext(f)[1][1:]} for f in files],
            "large_file": {"path": F["returns"], "rows": info["returns_rows"]},
            "distractor_files": sorted(DISTRACTORS)}
    with open(os.path.join(out, "metadata.json"), "w", encoding="utf-8") as f:
        json.dump(meta, f, indent=2, ensure_ascii=False)
        f.write("\n")
    WR.set_mtime(os.path.join(out, "metadata.json"), MTIME["index"])
    return meta


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    ap.add_argument("--record")
    a = ap.parse_args()
    T = os.path.join(a.out, "target")
    if os.path.isdir(T):
        shutil.rmtree(T)
    os.makedirs(T)
    W = core.make_world()
    info = write_data(W, T)
    L = load(T)
    R = compute(L)
    status, touched = write_padro(W, T, R)
    L["jun"] = A.read_csv(os.path.join(T, F["padro06"]))
    L["sep"] = A.read_csv(os.path.join(T, F["padro09"]))
    Dd = write_docs(W, T, L, R)
    WR.write_csv(os.path.join(T, F["deposits"]), distractors.DEPOSIT_HEADER,
                 distractors.deposits(W["stock"], W["nw_sections"]))
    lh_ids = {h for h in W["holders"] if not h.startswith("old")}
    WR.write_csv(os.path.join(T, F["visits"]), distractors.VISIT_HEADER,
                 distractors.visits(W["stock"], W["timeline"], {h: W["holders"][h]["nif"] for h in lh_ids}, lh_ids))
    write_index(T)
    finish_containers(T)
    meta = metadata(a.out, T, info)
    import asserts
    record = asserts.run(W, T, L, R, Dd, F, meta, info, touched)
    if a.record:
        with open(a.record, "w", encoding="utf-8") as f:
            json.dump(record, f, indent=1, sort_keys=True, default=str)
    print(f"task130 pack built: {len(os.listdir(T))} files, {record['n_assert']} assertions green")


if __name__ == "__main__":
    main()
