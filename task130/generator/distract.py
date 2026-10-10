"""Files the TA team keeps beside the extracts that the call-off does not use."""
import datetime as dt

import numpy as np
import pandas as pd
import openpyxl
import xlcalc
from openpyxl.styles import Font

from docs import _hdr

BANDS = ["Band A (urgent medical or welfare)", "Band B (overcrowded, two bedrooms short)", "Band C (other housing need)",
         "Transfer (existing social tenant)", "Sheltered and extra care"]


def register_lets_csv(out_path, rng):
    """Social lets in the borough April 2025 to March 2026 through the housing register and transfers.
    Move-on lets to households owed the main duty are recorded in the move-on lets log, not here."""
    rows = []
    days = pd.date_range("2025-04-01", "2026-03-31", freq="D")
    n = 812
    landlords = ["Medway Valley Homes", "Orchard Vale Housing Association", "Kentish Weald Homes", "Larkmoor Homes",
                 "Saxon Gate Housing Trust", "Downsview Housing"]
    areas = ["Larkmoor Town", "Wharfside", "Eastbrook", "Hollins Green", "St Mildred's", "Copperhill", "Saltings"]
    for k in range(n):
        band = rng.choice(BANDS, p=[0.18, 0.21, 0.27, 0.22, 0.12])
        beds = int(rng.choice([0, 1, 2, 3, 4], p=[0.06, 0.38, 0.33, 0.19, 0.04]))
        if band.startswith("Sheltered"):
            beds = int(rng.choice([0, 1], p=[0.4, 0.6]))
        rows.append(dict(let_date=days[int(rng.integers(0, len(days)))].strftime("%Y-%m-%d"),
                         band=band, bedrooms="bedsit" if beds == 0 else ("4+" if beds >= 4 else str(beds)),
                         landlord=landlords[int(rng.integers(0, len(landlords)))],
                         area=areas[int(rng.integers(0, len(areas)))],
                         weeks_on_register=int(max(1, rng.gamma(2.2, 38)))))
    df = pd.DataFrame(rows).sort_values(["let_date", "landlord"], kind="mergesort").reset_index(drop=True)
    df.insert(0, "let_id", ["HRL%04d" % (k + 1) for k in range(len(df))])
    df.to_csv(out_path, index=False)


def spend_xlsx(out_path):
    wb = openpyxl.Workbook()
    wb.properties.creator = "Debra Young"
    wb.properties.lastModifiedBy = "Debra Young"
    wb.properties.created = dt.datetime(2026, 7, 14, 11, 0)
    wb.properties.modified = dt.datetime(2026, 7, 21, 15, 45)
    wb.properties.title = "TA spend monitoring 2026-27 Q1"
    ws = wb.active
    ws.title = "Q1"
    ws["A1"] = "Temporary accommodation: revenue monitoring, quarter 1 2026/27 (April to June)"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "GBP thousands. Gross cost before housing benefit subsidy and client contributions."
    ws["A2"].font = Font(italic=True, size=9)
    _hdr(ws, 4, ["Line", "Annual budget", "Q1 budget", "Q1 actual", "Q1 variance", "Forecast outturn", "Comment"],
         [34, 14, 12, 12, 12, 16, 60])
    lines = [
        ("Bed and breakfast", 3120, 780, 846, "Families over six weeks in B&B: 11 at 30 June."),
        ("Hostels (council-run)", 1410, 352, 349, ""),
        ("Nightly-paid self-contained (Lot 2)", 6880, 1720, 1897, "Volume above plan; rates as schedule."),
        ("Private sector leases", 4215, 1054, 1041, ""),
        ("Tenter House block lease", 2530, 633, 633, "Fixed rent, 178 flats."),
        ("Housing benefit subsidy and charges", -9760, -2440, -2398, "Subsidy capped at 90 per cent of January 2011 LHA."),
    ]
    r = 5
    for name, ann, qb, qa, com in lines:
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=ann)
        ws.cell(row=r, column=3, value=qb)
        ws.cell(row=r, column=4, value=qa)
        ws.cell(row=r, column=5, value="=D%d-C%d" % (r, r))
        ws.cell(row=r, column=6, value=round(ann + (qa - qb) * 3.1))
        ws.cell(row=r, column=7, value=com)
        r += 1
    ws.cell(row=r, column=1, value="Net cost").font = Font(bold=True)
    for j in range(2, 7):
        col = openpyxl.utils.get_column_letter(j)
        ws.cell(row=r, column=j, value="=SUM(%s5:%s%d)" % (col, col, r - 1)).font = Font(bold=True)
    ws.freeze_panes = "A5"
    xlcalc.freeze(wb)
    wb.save(out_path)
