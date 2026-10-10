#!/usr/bin/env python3
"""Independent verifier for task130. Reads only the shipped files under <task>/target (and, for the
input gates alone, <task>/metadata.json). Shares no code with the generator: its own parsing, its own
conformance, its own roll, its own asks. Recomputes every rung, rival-killer, calibration outcome and
graded figure, and prints them; exits non-zero on any failed check.

    python3 -I verify.py /path/to/task130 [--json out.json]
"""
import argparse
import datetime as dt
import json
import os
import re
import sys
import zipfile
from decimal import Decimal, ROUND_HALF_UP

import pandas as pd
from openpyxl import load_workbook

N_CHECKS = [0]


def check(cond, msg):
    if not cond:
        print("FAIL:", msg)
        sys.exit(1)
    N_CHECKS[0] += 1


def csv(path):
    return pd.read_csv(path, dtype=str, keep_default_na=False)


def one_dec(x):
    return Decimal(x).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)


def pct(n, d):
    return Decimal(int(n)) * Decimal(100) / Decimal(int(d))


class Pack:
    def __init__(self, task):
        t = os.path.join(task, "target")
        self.t = t
        ls = sorted(os.listdir(t))
        pick = lambda pat: os.path.join(t, next(f for f in ls if re.search(pat, f)))
        self.files = ls
        self.ret = csv(pick(r"^declaracions_.*\.csv$"))
        self.cad = csv(pick(r"^cadastre_.*\.csv$"))
        self.deeds = csv(pick(r"^escriptures_.*\.csv$"))
        self.led = csv(pick(r"^EPPR_.*\.csv$"))
        self.jun = csv(pick(r"^padro_lliurament_2026-06\.csv$"))
        self.sep = csv(pick(r"^padro_lliurament_2026-09\.csv$"))
        self.codes = csv(pick(r"^codis_grup\.csv$"))
        wb = load_workbook(pick(r"^registre_.*\.xlsx$"), read_only=True)
        tit = list(wb["Titulars"].iter_rows(values_only=True))
        self.tit = pd.DataFrame(tit[1:], columns=tit[0]).fillna("").astype(str)
        lk = list(wb["Vinculacions"].iter_rows(values_only=True))
        self.links = pd.DataFrame(lk[1:], columns=lk[0]).fillna("").astype(str)
        ws = load_workbook(pick(r"^annex_designacio_2026\.xlsx$"), read_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        hi = next(i for i, r in enumerate(rows) if r[0] == "Secció censal")
        self.annex = pd.DataFrame(rows[hi + 1:], columns=rows[hi])
        self.bull_text = pdf_text(pick(r"2026T2\.pdf$"))
        res = self.cad[self.cad["us"] == "Residencial"]
        self.watch = sorted(set(self.annex["Secció censal"].astype(str)))
        self.N = res.groupby("seccio_censal").size().to_dict()
        self.sec = dict(zip(self.cad["referencia_cadastral"], self.cad["seccio_censal"]))
        self.unit = {(p, int(c)): r for r, p, c in zip(self.cad["referencia_cadastral"], self.cad["parcela"], self.cad["carrec"])}
        self.psec = dict(zip(self.cad["parcela"], self.cad["seccio_censal"]))
        self.pre18 = {}
        for r in self.cad["referencia_cadastral"]:
            self.pre18.setdefault(r[:18], []).append(r)
        self.reg = set(self.tit.loc[self.tit["data_baixa"] == "", "nif"])
        self.name = dict(zip(self.tit["nif"], self.tit["nom"]))
        self.gname = dict(zip(self.codes["codi_grup"], self.codes["nom_grup"]))
        self.galta = dict(zip(self.codes["codi_grup"], self.codes["data_alta"]))


def pdf_text(path):
    from pypdf import PdfReader
    return "\n".join(p.extract_text() for p in PdfReader(path).pages)


def fix_ref(P, raw):
    s = re.sub(r"[^0-9A-Za-zÑñ]", "", raw).upper()
    if s in P.sec:
        return s
    c = P.pre18.get(s[:18], [])
    return c[0] if len(c) == 1 else None


def units(text):
    out = []
    for part in text.split(","):
        a, _, b = part.partition("-")
        out.extend(range(int(a), int(b or a) + 1))
    return out


def lodgements(P, cutoff):
    df = P.ret[P.ret["data_presentacio"] <= cutoff]
    return {k: g for k, g in df.groupby(["nif_declarant", "trimestre"], sort=True)}


def effective(g, versions=True):
    """Rows standing for one holder-quarter group."""
    order = g.drop_duplicates("id_declaracio").sort_values(["data_presentacio", "id_declaracio"])
    if not versions:
        return g[g["id_declaracio"] == order.iloc[-1]["id_declaracio"]]
    keep = []
    for _, v in order.iterrows():
        keep = [v["id_declaracio"]] if v["tipus_declaracio"] in ("O", "S") else keep + [v["id_declaracio"]]
    return g[g["id_declaracio"].isin(keep)]


def picture(P, quarter, cutoff, rung):
    """rung 0, 1 or 2. Returns list of (key, holder, section, sched) where sched=(date, buyer) or None."""
    L = lodgements(P, cutoff)
    holders = sorted({h for h, _ in L})
    out = []
    for h in holders:
        if rung == 0:
            if (h, quarter) not in L:
                continue
            rows = effective(L[(h, quarter)], versions=False)
        else:
            qs = sorted(q for (hh, q) in L if hh == h and q <= quarter)
            if not qs:
                continue
            rows = effective(L[(h, qs[-1])])
        for _, r in rows.iterrows():
            sch = (r["data_transmissio_prevista"], r["nif_adquirent_previst"]) if r["data_transmissio_prevista"] else None
            if rung < 2:
                if r["unitats"]:
                    if rung == 0 or r["referencia_cadastral"] not in P.psec:
                        continue
                    out.append((r["referencia_cadastral"] + "|" + r["id_declaracio"], h, P.psec[r["referencia_cadastral"]], sch))
                elif r["referencia_cadastral"] in P.sec:
                    out.append((r["referencia_cadastral"] + ("" if r["codi_tinenca"] == "PD" else "|x" + h), h,
                                P.sec[r["referencia_cadastral"]], sch))
                continue
            if r["codi_tinenca"] != "PD":
                continue
            if r["unitats"]:
                parcel = re.sub(r"\s", "", r["referencia_cadastral"]).upper()
                for u in units(r["unitats"]):
                    ref = P.unit[(parcel, u)]
                    out.append((ref, h, P.sec[ref], sch))
            else:
                ref = fix_ref(P, r["referencia_cadastral"])
                check(ref is not None, f"reference {r['referencia_cadastral']} resolves")
                out.append((ref, h, P.sec[ref], sch))
    return out


def tally(items, P, watch):
    c = {s: 0 for s in watch}
    for k, h, s, _ in items:
        if s in c:
            c[s] += 1
    return c


def commitments(P, interval=120, fn=None):
    out = {}
    d = P.led[(P.led["programa"] == "PPO") & (P.led["fase"] == "D")]
    for day, text in zip(d["data_comptable"], d["concepte"]):
        if "tempteig" not in text:
            continue
        ref = text.split("RC ")[1].split(".")[0].strip()
        c = dt.date.fromisoformat(day)
        comp = fn(c) if fn else c + dt.timedelta(days=interval)
        planned = text.rsplit("prevista ", 1)[1].strip()
        out[ref] = (c, comp, dt.date(int(planned[6:10]), int(planned[3:5]), int(planned[0:2])))
    return out


def roll(P, items, start, deed_end, target, mode="R3", tk=None):
    own = {k: h for k, h, s, _ in items}
    sch = {k: x for k, h, s, x in items if x}
    d = P.deeds[(P.deeds["data_atorgament"] > start) & (P.deeds["data_atorgament"] <= deed_end)
                & (P.deeds["data_inscripcio"] <= deed_end)].sort_values(["data_atorgament", "num_entrada"])
    for ref, buyer in zip(d["referencia_cadastral"], d["nif_adquirent"]):
        if buyer in P.reg:
            own[ref] = buyer
        else:
            own.pop(ref, None)
    for ref, (day, buyer) in sorted(sch.items()):
        if ref not in own or day <= deed_end:
            continue
        if mode != "R3" and tk and ref in tk:
            c, comp, planned = tk[ref]
            when = {"R4": day, "R4p": c.isoformat(), "R5": comp.isoformat()}[mode]
            if when <= target:
                own.pop(ref)
            continue
        if day <= target:
            if buyer in P.reg:
                own[ref] = buyer
            else:
                own.pop(ref)
    return own


def per_section(P, own, watch):
    c = {s: 0 for s in watch}
    for k in own:
        s = P.sec.get(k.split("|")[0])
        if s in c:
            c[s] += 1
    return c


def chosen(P, counts, watch):
    return sorted(s for s in watch if pct(counts[s], P.N[s]) >= 25)


def members(P, day):
    m = {}
    for _, r in P.links.iterrows():
        if r["data_efecte_inici"] <= day and (r["data_efecte_fi"] == "" or day <= r["data_efecte_fi"]):
            m[r["nif"]] = r["codi_grup"]
    return m


def biggest(P, own, watch, mem):
    out = {}
    for s in watch:
        c = {}
        for k, h in own.items():
            if P.sec.get(k) == s:
                g = mem.get(h)
                key = ("g", g) if g else ("h", h)
                c[key] = c.get(key, 0) + 1
        r = sorted(c.items(), key=lambda kv: (-kv[1], kv[0]))
        name = P.gname[r[0][0][1]] if r[0][0][0] == "g" else P.name[r[0][0][1]]
        out[s] = (name, r[0][1], r[1][1])
    return out


def lapsed(doc, alta, renov, ref_day):
    if doc not in ("TIE-T", "PAS"):
        return False
    base = dt.date.fromisoformat(renov or alta)
    try:
        dl = base.replace(year=base.year + 2)
    except ValueError:
        dl = base.replace(year=base.year + 2, day=28)
    return dl < ref_day


def occupied(P):
    sep_d = set(zip(P.sep["codi_municipi"], P.sep["districte"]))
    occ = set()
    for df, ref_day, use in ((P.sep, dt.date(2026, 9, 1), lambda k: True),
                             (P.jun, dt.date(2026, 6, 1), lambda k: k not in sep_d)):
        for _, r in df.iterrows():
            if not use((r["codi_municipi"], r["districte"])):
                continue
            if r["persones"] != "":
                if int(r["persones"]) > 0:
                    occ.add(r["referencia_cadastral"])
            elif not lapsed(r["tipus_document"], r["data_alta"], r["data_ultima_renovacio"], ref_day):
                occ.add(r["referencia_cadastral"])
    return occ


def bulletin_figures(text, watch):
    """Rows print the section as municipality code, district and section; the 10-digit code chains them."""
    flat = " ".join(text.split())
    out = {}
    for m in re.finditer(r"(?<!\d)(\d{5}) (\d{2}) (\d{3}) (\d[\d.]*) (\d[\d.]*) (\d+,\d)", flat):
        code = m.group(1) + m.group(2) + m.group(3)
        if code in watch:
            out[code] = (int(m.group(4).replace(".", "")), int(m.group(5).replace(".", "")))
    check(set(out) == set(watch), f"bulletin rows parsed for {len(out)} of {len(watch)} sections")
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("task")
    ap.add_argument("--json")
    a = ap.parse_args()
    P = Pack(a.task)
    W = P.watch
    check(len(W) == 14, "fourteen watch-list sections in the 2026 annex")
    pics = {r: picture(P, "2026T2", "2026-08-31", r) for r in (0, 1, 2)}
    C = {f"R{r}": tally(pics[r], P, W) for r in (0, 1, 2)}
    tk = commitments(P)
    rolled = {}
    for m in ("R3", "R4", "R4p", "R5"):
        rolled[m] = roll(P, pics[2], "2026-06-30", "2026-09-30", "2027-01-01", m, tk)
        C[m] = per_section(P, rolled[m], W)
    sets = {m: chosen(P, C[m], W) for m in C}
    check(len({tuple(v) for v in sets.values()}) == 7, "seven rungs, seven distinct sets")
    ans = sets["R5"]
    check(len(ans) == 6, "the answer designates six sections")
    # controls
    bull = bulletin_figures(P.bull_text, W)
    check(all(bull[s] == (P.N[s], C["R2"][s]) for s in W), "R2 reproduces the 30 June bulletin on 14 of 14")
    check(sum(abs(C["R1"][s] - C["R2"][s]) >= 4 for s in W) >= 8, "R1 misses the bulletin on 8 or more sections")
    check(sum(abs(C["R0"][s] - C["R2"][s]) >= 4 for s in W) >= 11, "R0 misses the bulletin on 11 or more sections")
    pic25 = picture(P, "2025T2", "2025-08-31", 2)
    r25 = per_section(P, roll(P, pic25, "2025-06-30", "2025-09-30", "2026-01-01"), W)
    ann = {str(r["Secció censal"]): int(r["Habitatges de grans tenidors"]) for _, r in P.annex.iterrows()}
    check(all(r25[s] == ann[s] for s in W), "the forward roll on the 2025 data reproduces the 2026 annex on 14 of 14")
    check(sum(tally(pic25, P, W)[s] != ann[s] for s in W) >= 9, "the 30 June 2025 snapshot misses the 2026 annex on 9+")
    # settled purchases and the clock
    settled = []
    for ref, (c, comp, planned) in tk.items():
        later = P.deeds[(P.deeds["referencia_cadastral"] == ref) & (P.deeds["data_atorgament"] > c.isoformat())]
        if len(later):
            settled.append((ref, c, dt.date.fromisoformat(later.iloc[0]["data_atorgament"]), planned))
    check(len(settled) == 23, "23 settled first-offer purchases")
    gaps = sorted({(deed - c).days for _, c, deed, _ in settled})
    check(gaps == [120], f"deed minus commitment {gaps}")
    check(min(abs((deed - pl).days) for _, c, deed, pl in settled) >= 9, "the notified date misses every settled deed")
    # back-test of agreed sales to 30 September
    agreed = {}
    for _, r in P.ret[(P.ret["data_transmissio_prevista"] != "") & (P.ret["data_transmissio_prevista"] <= "2026-09-30")].iterrows():
        refs = ([P.unit[(re.sub(r"\s", "", r["referencia_cadastral"]).upper(), u)] for u in units(r["unitats"])]
                if r["unitats"] else [fix_ref(P, r["referencia_cadastral"])])
        for ref in refs:
            agreed[(ref, r["data_transmissio_prevista"], r["nif_adquirent_previst"])] = 1
    dk = set(zip(P.deeds["referencia_cadastral"], P.deeds["data_atorgament"], P.deeds["nif_adquirent"]))
    check(all(k in dk for k in agreed), f"back-test: {len(agreed)} of {len(agreed)} agreed sales completed as agreed")
    # corridor
    for n in range(110, 133):
        t2 = commitments(P, n)
        assert per_section(P, roll(P, pics[2], "2026-06-30", "2026-09-30", "2027-01-01", "R5", t2), W) == C["R5"], n
    check(True, "corridor 110 to 132 days")
    # asks
    mem = members(P, "2027-01-01")
    big = biggest(P, rolled["R5"], W, mem)
    check(all(big[s][1] - big[s][2] >= 3 for s in W), "largest group leads by 3 or more")
    for m in ("R3", "R4", "R4p"):
        check(biggest(P, rolled[m], W, mem) == big, f"ask B identical under {m}")
    occ = occupied(P)
    empty = {s: sum(1 for k in rolled["R5"] if P.sec.get(k) == s and k not in occ) for s in W}
    for m in ("R3", "R4", "R4p"):
        e2 = {s: sum(1 for k in rolled[m] if P.sec.get(k) == s and k not in occ) for s in W}
        check(e2 == empty, f"ask C identical under {m}")
    sh = {s: pct(C["R5"][s], P.N[s]) for s in W}
    des = [s for s in W if sh[s] >= 25]
    und = [s for s in W if sh[s] < 25]
    nd, nu = min(des, key=lambda s: sh[s]), max(und, key=lambda s: sh[s])
    check(all(abs(sh[s] - 25) >= Decimal("0.5") for s in W), "every share 0.5 points or more from the line")
    # bridge
    bridge = {}
    for s in (nu, nd):
        prev, steps = C["R2"][s], []
        start = prev
        for mth in range(7, 13):
            end = (dt.date(2026, mth + 1, 1) - dt.timedelta(days=1)) if mth < 12 else dt.date(2026, 12, 31)
            de = min(end, dt.date(2026, 9, 30)).isoformat()
            c = per_section(P, roll(P, pics[2], "2026-06-30", de, end.isoformat(), "R5", tk), W)[s]
            steps.append(c - prev)
            prev = c
        check(start + sum(steps) == C["R5"][s], f"bridge {s} closes")
        bridge[s] = (start, steps)
    # input gates
    meta = json.load(open(os.path.join(a.task, "metadata.json"), encoding="utf-8"))
    check(len(P.files) >= 10 and len({os.path.splitext(f)[1] for f in P.files}) >= 3, "input gates: files and formats")
    check(len(P.ret) >= 25000, f"input gate: {len(P.ret)} return rows")
    check(len(meta["distractor_files"]) >= 2 and set(meta["distractor_files"]) <= set(P.files), "distractors present")
    out = {"answer": ans, "sets": sets, "counts": C,
           "golden": {s: {"lh": C["R5"][s], "share": str(one_dec(sh[s])), "group": big[s][0], "group_n": big[s][1],
                          "empty": empty[s]} for s in W},
           "nearest": {"designated": [nd, str(one_dec(sh[nd])), str(one_dec(sh[nd] - 25))],
                       "undesignated": [nu, str(one_dec(sh[nu])), str(one_dec(25 - sh[nu]))]},
           "bridge": bridge, "settled": len(settled), "backtest": len(agreed), "annex2026": r25,
           "checks": N_CHECKS[0]}
    print(f"verify: {N_CHECKS[0]} checks green; answer {ans}; nearest designated {out['nearest']['designated']}, "
          f"undesignated {out['nearest']['undesignated']}")
    if a.json:
        with open(a.json, "w", encoding="utf-8") as f:
            json.dump(out, f, indent=1, sort_keys=True, default=str)


if __name__ == "__main__":
    main()
