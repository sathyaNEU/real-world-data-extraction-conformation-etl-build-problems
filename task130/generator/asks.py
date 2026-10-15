"""The two asks from the shipped pack, with the stops a solver lands on when a device or hazard is mishandled.

Ask A: for each entity, whether any of the open claims under its policies with its expiring carrier is missing from
that carrier's latest loss run and, if so, how many and the register's reserves on them.
Ask B: the covered clinician FTE the payroll paid at each entity in July, August and September 2026, from the HCM pay
lines."""
from __future__ import annotations

import csv
import datetime as dt
import os

import openpyxl

import common as K

QM_NAMES = {"Tolliver Falls Hospital": "TFH", "Alderwick Medical Center": "AMC"}
SC_NAMES = {"CORVANE MEDICAL GROUP": "CMG", "BRENNICK COMMUNITY HOSP": "BCH"}
ASK_B_MONTHS = ("2026-07", "2026-08", "2026-09")
TRAINEE_CODES = ("GME-RES", "GME-FEL")
HOURS_FULL = 173.33


def _d(s):
    return dt.date.fromisoformat(s) if s else None


def _f(s):
    return float(s) if s not in ("", None) else None


def norm(ref):
    return "".join(ch for ch in str(ref).upper() if ch.isalnum())


def load_register(target):
    rows = []
    with open(os.path.join(target, K.REGISTER), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            rows.append(dict(claim_no=r["claim_no"], incident_no=r["incident_no"], defendant_seq=int(r["defendant_seq"]),
                             clinician_id=r["clinician_id"], defendant_role=r["defendant_role"],
                             specialty=r["specialty"], coverage_line=r["coverage_line"],
                             entity_at_incident=r["entity_at_incident"], incident_date=_d(r["incident_date"]),
                             claim_made_date=_d(r["claim_made_date"]), matter_type=r["matter_type"],
                             allegation=r["allegation"], status=r["status"], closed_date=_d(r["closed_date"]),
                             indemnity_paid=_f(r["indemnity_paid"]), defense_paid=_f(r["defense_paid"]),
                             indemnity_reserve=_f(r["indemnity_reserve"]), defense_reserve=_f(r["defense_reserve"]),
                             prior_claim_no=r["prior_claim_no"]))
    return rows


def load_qm(target, fname):
    wb = openpyxl.load_workbook(os.path.join(target, fname), data_only=True)
    ws = wb["Loss Run"]
    valued = None
    for row in ws.iter_rows(min_row=1, max_row=6, values_only=True):
        if row[0] and str(row[0]).startswith("Valuation Date"):
            m, d, y = str(row[0]).split(": ")[1].split("/")
            valued = dt.date(int(y), int(m), int(d))
    out = []
    for row in ws.iter_rows(min_row=8, values_only=True):
        if not row[3]:
            continue
        out.append(dict(entity=QM_NAMES[row[2]], ref=row[4], ctype=row[7], status=row[8]))
    return valued, out


def load_sc(target):
    out = []
    valued = None
    with open(os.path.join(target, K.SC_SEP), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            out.append(dict(entity=SC_NAMES[r["INSURED_NAME"]], ref=r["INSURED_REF"], ctype=r["CLM_TYPE"],
                            status=r["CLM_STATUS"]))
            m, d, y = r["VALUATION_DT"].split("/")
            valued = dt.date(int(y), int(m), int(d))
    return valued, out


def load_notices(target):
    with open(os.path.join(target, K.NOTICES), encoding="utf-8") as fh:
        return {r["notice_no"]: r["incident_no"] for r in csv.DictReader(fh)}


def carrier_terms(target):
    """{entity: [(start, end, carrier)]} from the premium history, in term order."""
    wb = openpyxl.load_workbook(os.path.join(target, K.PREMIUMS), data_only=True)
    ws = wb["Premium history"]
    names = {v["name"]: k for k, v in K.ENTITIES.items()}
    out = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not row[0]:
            continue
        start = row[3].date() if hasattr(row[3], "date") else row[3]
        end = row[4].date() if hasattr(row[4], "date") else row[4]
        out.setdefault(names[row[0]], []).append((start, end, row[1]))
    for v in out.values():
        v.sort()
    return out


def expiring_era_start(terms):
    """The first day of the unbroken run of terms the expiring carrier (the last term's carrier) wrote."""
    car = terms[-1][2]
    start = terms[-1][0]
    for a, b, c in reversed(terms):
        if c != car:
            break
        start = a
    return start, car


def ask_a(target, rows=None):
    rows = rows if rows is not None else load_register(target)
    notices = load_notices(target)
    claim_inc = {norm(r["claim_no"]): r["incident_no"] for r in rows}
    notice_inc = {norm(k): v for k, v in notices.items()}
    qm_v, qm = load_qm(target, K.QM_SEP)
    qmj_v, qmj = load_qm(target, K.QM_JUN)
    sc_v, sc = load_sc(target)
    assert qm_v == K.VALUATION and sc_v == K.VALUATION and qmj_v == dt.date(2026, 6, 30)
    latest = {"TFH": qm, "AMC": qm, "CMG": sc, "BCH": sc}
    june = {"TFH": qmj, "AMC": qmj, "CMG": sc, "BCH": sc}
    terms = carrier_terms(target)
    inc_date = {r["incident_no"]: r["incident_date"] for r in rows}
    no_inc = {r["claim_no"]: r["incident_no"] for r in rows}
    rev_notice = {}
    for k, v in notices.items():
        rev_notice.setdefault(v, set()).add(norm(k))

    def found(run, e, notices_too=True):
        out = set()
        for x in run:
            if x["entity"] != e:
                continue
            k = norm(x["ref"])
            if k in claim_inc:
                out.add(claim_inc[k])
            elif notices_too and k in notice_inc:
                out.add(notice_inc[k])
        return out

    def reserves_of(sel, incs):
        return sum((r["indemnity_reserve"] or 0.0) + (r["defense_reserve"] or 0.0)
                   for r in rows if sel(r) and r["incident_no"] in incs)

    res = {}
    for e in K.ENTITY_ORDER:
        line = K.ENTITIES[e]["line"]
        on_line = lambda r, line=line: r["coverage_line"] == line
        opn = sorted({r["incident_no"] for r in rows if on_line(r) and r["status"] in ("OPEN", "REOPENED")})
        era, car = expiring_era_start(terms[e])
        exp = [i for i in opn if inc_date[i] >= era]
        old = [i for i in opn if inc_date[i] < era]
        on = found(latest[e], e)
        miss = [i for i in exp if i not in on]
        out = dict(open=len(opn), expiring=len(exp), old_era=old, era_start=era, carrier=car, missing=miss,
                   count=len(miss), reserves=reserves_of(on_line, set(miss)))
        # S1: every open claim on the coverage line, the expiring carrier's era not applied
        s1 = [i for i in opn if i not in on]
        out["s1"] = (len(s1), reserves_of(on_line, set(s1)))
        # S2: the runs' notice references left unmatched
        on2 = found(latest[e], e, notices_too=False)
        s2 = [i for i in exp if i not in on2]
        out["s2"] = (len(s2), reserves_of(on_line, set(s2)))
        # S3: claim number by claim number, a suit not read through its prior claim number
        refs = {norm(x["ref"]) for x in latest[e] if x["entity"] == e}
        open_nos = sorted({r["claim_no"] for r in rows if on_line(r) and r["status"] in ("OPEN", "REOPENED")
                           and inc_date[r["incident_no"]] >= era})
        s3 = [c for c in open_nos if norm(c) not in refs and not (rev_notice.get(no_inc[c], set()) & refs)]
        out["s3"] = (len(s3), sum((r["indemnity_reserve"] or 0.0) + (r["defense_reserve"] or 0.0)
                                  for r in rows if on_line(r) and r["claim_no"] in s3))
        # S4: the run received last (Quillan's revision 1 of its June run) read as the latest
        onj = found(june[e], e)
        s4 = [i for i in exp if i not in onj]
        out["s4"] = (len(s4), reserves_of(on_line, set(s4)))
        # S5: the entity's open claims by entity_at_incident (claims naming only contracted clinicians included)
        opn5 = sorted({r["incident_no"] for r in rows if r["entity_at_incident"] == e and
                       r["status"] in ("OPEN", "REOPENED") and inc_date[r["incident_no"]] >= era})
        s5 = [i for i in opn5 if i not in on]
        out["s5"] = (len(s5), reserves_of(lambda r, e=e: r["entity_at_incident"] == e, set(s5)))
        res[e] = out
    return res


def crosswalk(target):
    with open(os.path.join(target, K.CROSSWALK), encoding="utf-8") as fh:
        return {r["job_profile_code"]: r["legacy_job_code"] for r in csv.DictReader(fh)}


def directory(target):
    with open(os.path.join(target, K.DIRECTORY), encoding="utf-8") as fh:
        return {r["cost_centre"]: r["entity_code"] for r in csv.DictReader(fh)}


def schedule_b():
    """Job code (either set) to class key, as Schedule B lists them."""
    out = {}
    for j in K.JOBS:
        if j[4]:
            out[j[1]] = j[4]
            out[j[2]] = j[4]
    return out


def hcm_lines(target):
    with open(os.path.join(target, K.HCM), encoding="utf-8") as fh:
        return [r for r in csv.DictReader(fh) if r["work_month"] in ASK_B_MONTHS]


def ask_b(target, lines=None):
    """Paid covered clinician FTE by entity and month: hours paid on the lines of each covered position (a Schedule B
    job through the crosswalk) over 173.33, on the entity of the line's cost centre, with a resident's or fellow's
    additional position left out (trainees are not Covered Clinicians)."""
    sb = schedule_b()
    xw = crosswalk(target)
    cc = directory(target)
    lines = lines if lines is not None else hcm_lines(target)
    trainee = {(r["worker_id"], r["work_month"]) for r in lines if xw[r["job_profile_code"]] in TRAINEE_CODES}
    names = ("gold", "s1", "s2", "s3", "s4", "s5")
    out = {n: {(e, m): 0.0 for e in K.ENTITY_ORDER for m in ASK_B_MONTHS} for n in names}
    seen = set()
    for r in lines:
        if not sb.get(xw[r["job_profile_code"]]):
            continue
        cell = (cc[r["cost_centre"]], r["work_month"])
        hrs = float(r["hours"])
        per_diem = float(r["fte"]) == 0.0 and hrs > 0
        out["s1"][cell] += hrs / HOURS_FULL                                 # trainees' additional positions kept
        if (r["worker_id"], r["work_month"]) in trainee:
            continue
        out["gold"][cell] += hrs / HOURS_FULL                               # paid hours, trainees left out
        if r["earning_code"] == "REG":
            out["s3"][cell] += hrs / HOURS_FULL                             # regular hours only
        if not per_diem:
            out["s5"][cell] += hrs / HOURS_FULL                             # per-diem clinicians left out
        out["s4"][cell] += float(r["fte"])                                  # the fte column on every line
        key = (r["worker_id"], r["position_id"], r["work_month"])
        if key not in seen:                                                 # the fte column, once a position
            seen.add(key)
            out["s2"][cell] += float(r["fte"])
    return out
