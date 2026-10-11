"""The two deliverables the prompt names, written from the shipped files: Tiffany Barrett's certification workbook and
Joyce Salinas's one-page board memo with its chart."""
import datetime as dt
import io
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import containers as C
import estimators as E
import world as W

BLUE, ORANGE = "#2a78d6", "#eb6834"
INK, INK2, GRID = "#1f1f1f", "#55534e", "#e3e2de"
T = (1, 2, 3, 4)


def whole(x):
    return int(round(x))


def figures(target):
    G, F = E.figures(target)
    out = {"n": G["n"], "f26": G["f26"], "base": G["base"], "prog": G["prog"], "prog_total": G["prog_total"],
           "an_detail": G["an_detail"], "f26_detail": G["f26_detail"], "R0": G["rung_tiers"]["R0"], "rungs": G["rungs"],
           "bid_n": G["bid_tiers"], "letter": G["letter"]}
    n = G["n"]
    out["R0_prog"] = (sum(n[t] * G["rung_tiers"]["R0"][t][0] for t in T) / 140,
                      sum(n[t] * G["rung_tiers"]["R0"][t][1] for t in T) / 140)
    apr = {}
    for t in T:
        apr[t] = sum(1 for m in G["members"]["bid"][t] if F.mm(m, dt.date(2025, 4, 1), dt.date(2025, 4, 30)) == 1)
    out["bid_apr"] = apr
    return out


# ------------------------------------------------------------------ the workbook
def workbook(path, R):
    wb = Workbook()
    thin = Side(style="thin", color="9E9E9E")
    hdr_fill = PatternFill("solid", fgColor="D9E2EC")
    tot_fill = PatternFill("solid", fgColor="F2F2F2")
    usd0, usd2, cnt = '"$"#,##0', '"$"#,##0.00', "#,##0"

    ws = wb.active
    ws.title = "Baseline"
    ws["A1"] = "Red Cedar Valley Schools Employee Health Plan"
    ws["A1"].font = Font(name="Calibri", size=13, bold=True)
    ws["A2"] = "Northcourt care management: baseline for the Programme Year, 1 April 2027 to 31 March 2028"
    ws["A2"].font = Font(name="Calibri", size=11, bold=True)
    ws["A3"] = "Certification under Schedule C.2 of Agreement RCV-2027-014. Prepared 12 March 2027 for the Board meeting of 25 March 2027."
    ws["A3"].font = Font(name="Calibri", size=10, italic=True, color="555555")
    heads = ["Risk tier", "Enrolled members", "2026 allowed PMPM, medical", "2026 allowed PMPM, pharmacy",
             "Baseline PMPM, medical", "Baseline PMPM, pharmacy"]
    r0 = 5
    for j, h in enumerate(heads, 1):
        c = ws.cell(row=r0, column=j, value=h)
        c.font = Font(name="Calibri", bold=True)
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=thin, top=thin)
    ws.row_dimensions[r0].height = 46
    for i, t in enumerate(T):
        r = r0 + 1 + i
        vals = ["Tier %d" % t, R["n"][t], whole(R["f26"][t][0]), whole(R["f26"][t][1]), whole(R["base"][t][0]), whole(R["base"][t][1])]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = Font(name="Calibri")
            if j == 2:
                c.number_format = cnt
            elif j > 2:
                c.number_format = usd0
            c.alignment = Alignment(horizontal="left" if j == 1 else "right")
    rp = r0 + 5
    for j in range(1, 8):
        ws.cell(row=rp - 1, column=j).border = Border(bottom=thin)
    hdr2 = {1: "Programme baseline PMPM, tiers weighted by enrolled members", 2: "Enrolled members", 5: "Medical",
            6: "Pharmacy", 7: "Total"}
    for j in range(1, 8):
        c = ws.cell(row=rp + 1, column=j, value=hdr2.get(j))
        c.font = Font(name="Calibri", bold=True)
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="left" if j == 1 else "center", vertical="center", wrap_text=True)
        c.border = Border(bottom=thin, top=thin)
    ws.row_dimensions[rp + 1].height = 32
    vals = {1: "All enrolled members", 2: sum(R["n"].values()), 5: whole(R["prog"][0]), 6: whole(R["prog"][1]),
            7: whole(R["prog_total"])}
    for j in range(1, 8):
        c = ws.cell(row=rp + 2, column=j, value=vals.get(j))
        c.font = Font(name="Calibri", bold=True)
        c.fill = tot_fill
        c.number_format = cnt if j == 2 else usd0
        c.alignment = Alignment(horizontal="left" if j == 1 else "right")
    notes = [
        "Whole dollars per member per month. No figure is rounded before the final figure (Schedule C.2), so a tier's medical and pharmacy "
        "figures are not added here.",
        "2026 allowed PMPM: the enrolled members' allowed claims incurred in calendar 2026 and paid through 28 February 2027, not "
        "completed, over their 2026 member months (Schedule C.2).",
        "Baseline PMPM: allowed claims per member month, 1 April 2025 to 31 March 2026, of the members Northcourt's model placed in each "
        "tier on 2024 claims in its RFP 2025-07 proposal (never enrolled: the award was suspended in April 2025), counted for the months "
        "each was enrolled, and projected 24 months at the renewal trend, 6.5% medical and 9.0% pharmacy (Schedule C.1(g)).",
        "Allowed claims include member cost sharing, are net of replacement and void claims and pharmacy reversals, before manufacturer "
        "rebates and stop-loss reimbursements (Schedule C.1(a)). Member month: enrolled on the first of the month (Schedule C.1(b)).",
        "Sources: Ellery claim lines and eligibility, Quarrystone claims (pulled 1 to 3 March 2027); Ellery conversion file of 1 January "
        "2025; Northcourt Final Enrolment List of 5 March 2027 and RFP 2025-07 stratification of 21 February 2025.",
    ]
    for i, ntx in enumerate(notes):
        c = ws.cell(row=rp + 4 + i, column=1, value=("Note %d. " % (i + 1)) + ntx)
        c.font = Font(name="Calibri", size=9, color="404040")
        ws.merge_cells(start_row=rp + 4 + i, start_column=1, end_row=rp + 4 + i, end_column=7)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[rp + 4 + i].height = 26 if len(ntx) < 160 else 38
    for j, wdt in enumerate([30, 11, 14, 14, 14, 14, 12], 1):
        ws.column_dimensions[get_column_letter(j)].width = wdt
    ws.freeze_panes = "A6"
    ws.print_area = "A1:G%d" % (rp + 4 + len(notes) - 1)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # working: the 2025 bid members
    w2 = wb.create_sheet("Bid members 2025-26")
    w2["A1"] = "Members Northcourt placed in each tier on 2024 claims (RFP 2025-07, 21 February 2025), April 2025 to March 2026"
    w2["A1"].font = Font(name="Calibri", bold=True)
    w2["A2"] = "Bramwell member numbers linked to Ellery member IDs through the conversion file; claims and months from 1 April 2025 to 31 March 2026."
    w2["A2"].font = Font(name="Calibri", size=9, italic=True, color="555555")
    heads = ["Risk tier", "Members on bid list", "Enrolled on 1 April 2025", "Member months", "Medical allowed",
             "Pharmacy allowed", "Medical PMPM", "Pharmacy PMPM", "Trend factor, medical", "Trend factor, pharmacy",
             "Baseline PMPM, medical", "Baseline PMPM, pharmacy"]
    for j, h in enumerate(heads, 1):
        c = w2.cell(row=4, column=j, value=h)
        c.font = Font(name="Calibri", bold=True)
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    w2.row_dimensions[4].height = 45
    fm, fr = W.TREND_MED ** 2, W.TREND_RX ** 2
    tot = [0, 0, 0, 0.0, 0.0]
    for i, t in enumerate(T):
        nb, mm, med, rx = R["an_detail"][t]
        row = ["Tier %d" % t, nb, R["bid_apr"][t], mm, round(med, 2), round(rx, 2), round(med / mm, 2), round(rx / mm, 2),
               fm, fr, round(R["base"][t][0], 2), round(R["base"][t][1], 2)]
        tot = [tot[0] + nb, tot[1] + R["bid_apr"][t], tot[2] + mm, tot[3] + med, tot[4] + rx]
        for j, v in enumerate(row, 1):
            c = w2.cell(row=5 + i, column=j, value=v)
            c.number_format = cnt if j in (2, 3, 4) else usd2 if j in (5, 6, 7, 8, 11, 12) else "0.000000" if j in (9, 10) else "General"
    trow = ["All tiers", tot[0], tot[1], tot[2], round(tot[3], 2), round(tot[4], 2)]
    for j, v in enumerate(trow, 1):
        c = w2.cell(row=9, column=j, value=v)
        c.font = Font(name="Calibri", bold=True)
        c.border = Border(top=thin)
        c.number_format = cnt if j in (2, 3, 4) else usd2 if j in (5, 6) else "General"
    w2["A11"] = ("Trend: 1.065 and 1.09 a year (renewal letter of 26 February 2027) over the 24 months from the middle of the "
                 "bid members' year (1 October 2025) to the middle of the Programme Year (1 October 2027).")
    w2["A12"] = ("Members who left the plan before 1 April 2025 have no member months in the year and drop out; members who "
                 "left during it count for the months they were enrolled.")
    for a in ("A11", "A12"):
        w2[a].font = Font(name="Calibri", size=9, color="404040")
    for j, wdt in enumerate([10, 11, 11, 10, 15, 15, 12, 12, 11, 11, 13, 13], 1):
        w2.column_dimensions[get_column_letter(j)].width = wdt
    w2.page_setup.orientation = "landscape"
    w2.page_setup.fitToWidth = 1
    w2.page_setup.fitToHeight = 0
    w2.sheet_properties.pageSetUpPr.fitToPage = True
    w2.freeze_panes = "B5"

    # working: the enrolled members' 2026
    w3 = wb.create_sheet("Enrolled members 2026")
    w3["A1"] = "Members on Northcourt's Final Enrolment List (5 March 2027), calendar 2026, paid through 28 February 2027"
    w3["A1"].font = Font(name="Calibri", bold=True)
    heads = ["Risk tier", "Enrolled members", "2026 member months", "Medical allowed", "Pharmacy allowed", "Medical PMPM",
             "Pharmacy PMPM", "Own 2026 PMPM at trend, medical", "Own 2026 PMPM at trend, pharmacy"]
    for j, h in enumerate(heads, 1):
        c = w3.cell(row=3, column=j, value=h)
        c.font = Font(name="Calibri", bold=True)
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    w3.row_dimensions[3].height = 45
    for i, t in enumerate(T):
        n_, mm, med, rx = R["f26_detail"][t]
        row = ["Tier %d" % t, n_, mm, round(med, 2), round(rx, 2), round(med / mm, 2), round(rx / mm, 2),
               round(R["R0"][t][0], 2), round(R["R0"][t][1], 2)]
        for j, v in enumerate(row, 1):
            c = w3.cell(row=4 + i, column=j, value=v)
            c.number_format = cnt if j in (2, 3) else usd2 if j > 3 else "General"
    w3["A9"] = ("Columns H and I are for reference only and are not certified: each tier's own 2026 PMPM carried 15 months "
                "at trend, 1 July 2026 to 1 October 2027. Weighted by enrolled members they come to %s medical and %s pharmacy, "
                "%s in total."
                % ("${:,.0f}".format(R["R0_prog"][0]), "${:,.0f}".format(R["R0_prog"][1]),
                   "${:,.0f}".format(R["R0_prog"][0] + R["R0_prog"][1])))
    w3["A9"].font = Font(name="Calibri", size=9, color="404040")
    w3["A10"] = ("Northcourt chose these members because of their 2026 claims, so their 2026 level is not what they would be "
                 "expected to cost in the Programme Year; see the Baseline sheet, note 3.")
    w3["A10"].font = Font(name="Calibri", size=9, color="404040")
    w3["A11"] = ("Where a member's coverage began partway through a month, as a newborn's does, that month's claims are included "
                 "and the month is not a member month (Schedule C.1(a) and (b)).")
    w3["A11"].font = Font(name="Calibri", size=9, color="404040")
    for j, wdt in enumerate([10, 11, 11, 15, 15, 12, 12, 15, 15], 1):
        w3.column_dimensions[get_column_letter(j)].width = wdt
    w3.page_setup.orientation = "landscape"
    w3.page_setup.fitToWidth = 1
    w3.page_setup.fitToHeight = 0
    w3.sheet_properties.pageSetUpPr.fitToPage = True
    w3.freeze_panes = "B4"

    when_c = dt.datetime(2027, 3, 9, 13, 26, 4, tzinfo=C.CENTRAL)
    when_m = dt.datetime(2027, 3, 12, 15, 48, 37, tzinfo=C.CENTRAL)
    wb.properties.creator = "Tiffany Barrett"
    wb.save(path)
    C.repack(path, C.core_props("Tiffany Barrett", "Joyce Salinas", when_c, when_m),
             C.app_props_xlsx([s.title for s in wb.worksheets], "Red Cedar Valley Schools"))


# ------------------------------------------------------------------ the chart
def chart_png(R):
    plt.rcParams.update({"font.family": "Carlito", "font.size": 8.5, "axes.edgecolor": GRID, "axes.linewidth": 0.8,
                         "xtick.color": INK2, "ytick.color": INK2, "text.color": INK, "svg.hashsalt": "rcv"})
    fig, axes = plt.subplots(1, 2, figsize=(6.6, 2.75), dpi=220)
    for k, (ax, label) in enumerate(zip(axes, ("Medical", "Pharmacy"))):
        x = range(4)
        a26 = [R["f26"][t][k] for t in T]
        bas = [R["base"][t][k] for t in T]
        w = 0.34
        ax.bar([i - w / 2 - 0.01 for i in x], a26, width=w, color=BLUE, label="2026 allowed PMPM (paid through 28 Feb 2027)", zorder=3)
        ax.bar([i + w / 2 + 0.01 for i in x], bas, width=w, color=ORANGE, label="Baseline PMPM, 2027-28", zorder=3)
        top = max(a26 + bas)
        for i, v in enumerate(bas):
            ax.text(i + w / 2 + 0.01, v + top * 0.02, "${:,.0f}".format(v), ha="center", va="bottom", fontsize=7, color=INK,
                    fontweight="bold", bbox=dict(boxstyle="square,pad=0.12", facecolor="white", edgecolor="none"), zorder=4)
        ax.set_xticks(list(x))
        ax.set_xticklabels(["Tier %d" % t for t in T])
        ax.set_ylim(0, top * 1.14)
        ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, p: "${:,.0f}".format(v)))
        ax.grid(axis="y", color=GRID, linewidth=0.6, zorder=0)
        ax.set_axisbelow(True)
        for sp in ("top", "right", "left"):
            ax.spines[sp].set_visible(False)
        ax.tick_params(axis="both", length=0, labelsize=7.5)
        ax.set_title(label, loc="left", fontsize=9, fontweight="bold", color=INK, pad=4)
    axes[0].set_ylabel("Dollars per member per month", fontsize=7.5, color=INK2)
    h, l = axes[0].get_legend_handles_labels()
    fig.legend(h, l, loc="lower center", ncol=2, frameon=False, fontsize=7.5, bbox_to_anchor=(0.5, -0.01))
    fig.suptitle("Medical baselines sit under half of 2026 cost in every tier; pharmacy moves both ways", x=0.01, ha="left",
                 fontsize=9.5, fontweight="bold", color=INK)
    fig.tight_layout(rect=(0, 0.08, 1, 0.95), w_pad=2.2)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None}, facecolor="white")
    plt.close(fig)
    return buf.getvalue()


# ------------------------------------------------------------------ the memo
def _para(doc, text, size=10.5, bold=False, italic=False, after=4, before=0, color=None, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if align:
        p.alignment = align
    return p


def memo(path, R, png):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.85)
    sec.top_margin, sec.bottom_margin = Inches(0.6), Inches(0.55)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.0
    hp = sec.header.paragraphs[0]
    hr = hp.add_run("Red Cedar Valley Schools, ISD 2817  |  Risk Management and Benefits")
    hr.font.size = Pt(8.5)
    hr.font.color.rgb = RGBColor.from_string("666666")
    fp = sec.footer.paragraphs[0]
    fr = fp.add_run("School Board meeting of 25 March 2027, item 7.3. Not for distribution outside the District and Northcourt.")
    fr.font.size = Pt(7.5)
    fr.font.color.rgb = RGBColor.from_string("777777")

    _para(doc, "MEMORANDUM", size=13, bold=True, after=6)
    tbl = doc.add_table(rows=5, cols=2)
    tbl.autofit = False
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    grid = tbl._tbl.tblGrid
    for gc, wdt in zip(grid.findall(qn("w:gridCol")), (1100, 8500)):
        gc.set(qn("w:w"), str(wdt))
    rows = [("To", "School Board (Deborah Johnson, Chair)"),
            ("From", "Joyce Salinas, Director of Risk Management and Benefits"),
            ("Copy", "Sandra Lopez, Superintendent; Anthony Olsen, Chief Financial Officer"),
            ("Date", "12 March 2027"),
            ("Subject", "Certifying the Northcourt care management baseline for 1 April 2027 to 31 March 2028")]
    for i, (a, b) in enumerate(rows):
        c0, c1 = tbl.rows[i].cells
        c0.width, c1.width = Inches(0.8), Inches(6.0)
        r0 = c0.paragraphs[0].add_run(a + ":")
        r0.bold = True
        r0.font.size = Pt(10)
        r1 = c1.paragraphs[0].add_run(b)
        r1.font.size = Pt(10)
        for c in (c0, c1):
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
    line = doc.add_paragraph()
    line.paragraph_format.space_after = Pt(6)
    pPr = line._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "808080")):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    pPr.append(bdr)

    m, r = R["prog"]
    tot = R["prog_total"]
    _para(doc, ("I recommend that the Board approve a programme baseline of ${:,} per member per month for the 140 members "
                "Northcourt enrolls on 1 April: ${:,} for medical claims and ${:,} for pharmacy. Northcourt's 40 per cent share "
                "of savings will be measured against the tier baselines that make it up, under Schedule C of the agreement the "
                "Board approved on 9 February.").format(whole(tot), whole(m), whole(r)), after=6)
    _para(doc, ("These members are on the list because 2026 was an expensive year for them, and for most it was the year of an "
                "operation, a heart attack, pneumonia, a cancer course or a premature birth that will not repeat. Northcourt ran the same model "
                "on our 2024 claims for its 2025 proposal, and because the Board suspended that award none of those 140 members "
                "was ever enrolled. Their claims from April 2025 to March 2026 show what a member the model flags costs in the "
                "following year without the programme, so each tier's baseline is that group's allowed claims per member month, "
                "counted for the months they were covered and brought forward 24 months at the 6.5 per cent medical and 9.0 per "
                "cent pharmacy trend in the February renewal. The four tiers are weighted by the members Northcourt enrolled in "
                "each."), after=6)
    r0m, r0r = R["R0_prog"]
    _para(doc, ("Carrying the enrolled members' own 2026 claims forward at trend would have produced ${:,}, {:.1f} times the "
                "figure above, and we would have paid Northcourt a share of savings that were always going to "
                "happen.").format(whole(r0m + r0r), (r0m + r0r) / tot), after=6)
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.space_after = Pt(1)
    pic.add_run().add_picture(io.BytesIO(png), width=Inches(6.6))
    _para(doc, ("Per member per month. 2026: allowed claims incurred in 2026 and paid through 28 February 2027, not completed, "
                "as Schedule C.2 asks. Baseline: 2025 proposal members by tier, April 2025 to March 2026, at trend. Source: "
                "Ellery and Quarrystone claims; Northcourt lists of 21 February 2025 and 5 March 2027."),
          size=8, italic=True, color="555555", after=6)
    _para(doc, ("Allowed claims here include member cost sharing and are taken before pharmacy rebates and stop-loss "
                "reimbursements, as Schedule C.1 defines them. Tier figures and the working are in cohort_baseline_2027.xlsx. "
                "Once the Board approves, I will send the certification to Northcourt, who have 30 days to dispute it under "
                "Schedule C.6."), after=0)
    doc.core_properties.author = "Joyce Salinas"
    doc.core_properties.title = "Northcourt baseline certification"
    doc.save(path)
    body = " ".join(p.text for p in doc.paragraphs)
    words = len(body.split())
    created = dt.datetime(2027, 3, 11, 8, 55, 12, tzinfo=C.CENTRAL)
    modified = dt.datetime(2027, 3, 12, 16, 22, 40, tzinfo=C.CENTRAL)
    C.repack(path, C.core_props("Joyce Salinas", "Joyce Salinas", created, modified, title="Northcourt baseline certification",
                                revision=6),
             C.app_props_docx("Red Cedar Valley Schools", 1, words, len(body)), drop=("docProps/thumbnail.jpeg",))


def write(out, target, S=None, L=None):
    gdir = os.path.join(out, "golden")
    os.makedirs(gdir, exist_ok=True)
    R = figures(target)
    workbook(os.path.join(gdir, "cohort_baseline_2027.xlsx"), R)
    png = chart_png(R)
    memo(os.path.join(gdir, "baseline_certification_memo.docx"), R, png)
    return R
