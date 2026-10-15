"""The goldens, computed from the pack in target/: trust_funding_2027.xlsx (the contribution and its exposure, the
carrier loss run check, the paid covered FTE for July to September 2026, notes) and trustees_paper_2027.pptx (the
contribution, then the chart of each entity's paid covered FTE in those three months), plus generator/figures.json
for the write-up."""
from __future__ import annotations

import io
import json
import math
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt     # noqa: E402
import matplotlib.ticker            # noqa: E402,F401
import openpyxl                     # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill   # noqa: E402
from openpyxl.utils import get_column_letter                # noqa: E402
from pptx import Presentation       # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.text import PP_ALIGN  # noqa: E402
from pptx.util import Cm, Pt         # noqa: E402

import analysis as AN    # noqa: E402
import asks as A         # noqa: E402
import calc as CA        # noqa: E402
import common as K       # noqa: E402
import postprocess as PP  # noqa: E402

BOOK, DECK = K.DELIVERABLES
MONTH_LABEL = {"2026-07": "July 2026", "2026-08": "August 2026", "2026-09": "September 2026"}
MONTH_SHORT = {"2026-07": "Jul", "2026-08": "Aug", "2026-09": "Sep"}
NAVY, SLATE, PAPER, LINE = "#1f2f3f", "#5a6672", "#ffffff", "#d5dbe1"
MONTH_COLOURS = {"2026-07": "#9db7cf", "2026-08": "#5f88ae", "2026-09": "#24476b"}
LAG_BANDS = [(0, 30, "Within 30 days"), (31, 90, "31 to 90 days"), (91, 364, "91 days to a year"),
             (365, 100000, "A year or more")]


def r1000(x):
    return int(math.floor(x / 1000.0 + 0.5) * 1000)


def r1(x):
    return math.floor(x * 10 + 0.5) / 10


def money(x):
    return "{:,}".format(r1000(x))


def long_date(d):
    return "%d %s %d" % (d.day, d.strftime("%B"), d.year)


def compute(target):
    m = AN.main_call(target)
    a = A.ask_a(target, m["rows"])
    b = A.ask_b(target)
    return dict(m=m, a=a, b=b)


def months_in(e):
    return 13 - K.ENTITIES[e]["transition"].month


def lag_profile(m):
    out = []
    for lo, hi, label in LAG_BANDS:
        cl = [c for c in m["ex"] if lo <= c["lag"] <= hi]
        out.append((label, len(cl), sum(c["cost"] for c in cl) / len(cl) if cl else 0.0,
                    sum(c["cost"] for c in cl)))
    return out


def figures(c):
    m, a, b = c["m"], c["a"], c["b"]
    fig = {"call": {"contribution": "{:,}".format(m["rounded"]), "contribution_usd": "$" + "{:,}".format(m["rounded"])}}
    fig["components"] = {"experience_claims": "{:,}".format(m["exp_claims"]),
                         "experience_fte_years": "{:,.1f}".format(r1(m["exp_fte_years"]))}
    fig["entities"] = {e: {"name": K.ENTITIES[e]["name"], "transition": long_date(K.ENTITIES[e]["transition"]),
                           "months_in_2027": months_in(e)} for e in K.ENTITY_ORDER}
    fig["ask_a"] = {}
    for e in K.ENTITY_ORDER:
        v = a[e]
        fig["ask_a"][e] = {"name": K.ENTITIES[e]["name"], "any_missing": "yes" if v["count"] else "no",
                           "count": v["count"], "reserves": money(v["reserves"]), "carrier": v["carrier"],
                           "carrier_since": long_date(v["era_start"])}
    fig["ask_b"] = {e: {m_: "%.1f" % r1(b["gold"][(e, m_)]) for m_ in A.ASK_B_MONTHS} for e in K.ENTITY_ORDER}
    fig["ask_b_total"] = {m_: "%.1f" % r1(sum(b["gold"][(e, m_)] for e in K.ENTITY_ORDER)) for m_ in A.ASK_B_MONTHS}
    fig["conventions"] = {"contribution_rounded_to": "{:,}".format(CA.BIN), "trend_pct": "5.0",
                          "hours_per_fte_month": "%.2f" % A.HOURS_FULL}
    fig["routes"] = {k: "{:,}".format(CA.rounded(v)) for k, v in m["routes"].items()}
    return fig


# ------------------------------------------------------------------------------------------------ workbook
HEAD = PatternFill("solid", fgColor="DCE3EA")
TOTAL = PatternFill("solid", fgColor="F1F3F5")
B = Font(name="Calibri", size=11, bold=True)
T = Font(name="Calibri", size=14, bold=True)
NOTE_FONT = Font(name="Calibri", size=10, italic=True, color="595959")


def _head(ws, row, values):
    for c, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=c, value=v if v != "" else None)
        cell.font = B
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal="left" if c == 1 else "right", wrap_text=True, vertical="bottom")


def write_book(path, c):
    m, a, b = c["m"], c["a"], c["b"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contribution"
    ws["A1"] = "Corvane Professional Liability Trust: 2027 contribution"
    ws["A1"].font = T
    ws["A2"] = "Expected cost of the Claims first made against the Trust in 2027, at 2027 cost level (funding policy 3.1)"
    ws["A4"] = "2027 Contribution ($)"
    ws["A4"].font = B
    ws["B4"] = m["rounded"]
    ws["B4"].font = B
    ws["B4"].number_format = "#,##0"
    ws["C4"] = "to the nearest $250,000"
    ws["C4"].font = NOTE_FONT
    ws["A6"] = "Entities in the Trust, 2027"
    ws["A6"].font = B
    _head(ws, 7, ["Participating Entity", "Transition Date", "Months in the Trust in 2027"])
    r = 8
    for e in K.ENTITY_ORDER:
        ws.cell(row=r, column=1, value=K.ENTITIES[e]["name"])
        ws.cell(row=r, column=2, value=long_date(K.ENTITIES[e]["transition"])).alignment = Alignment(horizontal="right")
        ws.cell(row=r, column=3, value=months_in(e))
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Experience").font = B
    r += 1
    for label, val, fmt in (("Claims from 2016 to 2024 incidents naming Covered Clinicians", m["exp_claims"], "#,##0"),
                            ("Covered FTE-years, 2016 to 2024 payroll", r1(m["exp_fte_years"]), "#,##0.0")):
        ws.cell(row=r, column=1, value=label)
        cell = ws.cell(row=r, column=3, value=val)
        cell.number_format = fmt
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Experience Claims by delay from incident to first making").font = B
    r += 1
    _head(ws, r, ["Delay", "Claims", "Average cost, 2027 level ($)", "Share of the experience's cost"])
    r += 1
    prof = lag_profile(m)
    tot_cost = sum(x[3] for x in prof)
    for label, n, avg, cost in prof:
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=n).number_format = "#,##0"
        ws.cell(row=r, column=3, value=r1000(avg)).number_format = "#,##0"
        cell = ws.cell(row=r, column=4, value=cost / tot_cost)
        cell.number_format = "0.0%"
        r += 1
    r += 1
    for t in ("Only Claims made by 31 December 2027 from incidents on or after each entity's Transition Date fall in "
              "2027. In the experience the Claims made within weeks of the care are disclosed harm settled through the "
              "Early Resolution Program, those made in the following months mostly small program requests, and the "
              "cost sits in the demands and suits made a year or more after the care.",
              "Source: risk management claims register extract valued 30 September 2026; legacy payroll 2016 to 2024; "
              "clinician staffing plan 2027, version 2026-09-29. Definitions on the Notes sheet."):
        ws.cell(row=r, column=1, value=t).font = NOTE_FONT
        r += 1
    for col, wd in zip(range(1, 5), [58, 18, 30, 30]):
        ws.column_dimensions[get_column_letter(col)].width = wd

    ws3 = wb.create_sheet("Carrier loss runs")
    ws3["A1"] = "Open claims under each entity's policies with its expiring carrier missing from that carrier's latest " \
                "loss run"
    ws3["A1"].font = T
    _head(ws3, 3, ["Participating Entity", "Expiring carrier", "Carrier's policies since", "Latest loss run",
                   "Open claims under its policies", "Missing from the loss run", "Our reserves on them ($)"])
    for i, e in enumerate(K.ENTITY_ORDER):
        v = a[e]
        vals = [K.ENTITIES[e]["name"], v["carrier"], long_date(v["era_start"]), "valued 30 September 2026",
                v["expiring"], v["count"] if v["count"] else "None", r1000(v["reserves"])]
        for col, val in enumerate(vals, start=1):
            cell = ws3.cell(row=4 + i, column=col, value=val)
            if col == 7:
                cell.number_format = "#,##0"
            if col >= 5:
                cell.alignment = Alignment(horizontal="right")
    rr = 10
    ws3.cell(row=rr - 1, column=1, value="Claims missing from the loss runs").font = B
    _head(ws3, rr, ["Participating Entity", "Claim number", "Incident number", "Incident date", "Claim made",
                    "Matter type", "Our reserves ($)"])
    rr += 1
    rows = m["rows"]
    n_missing = 0
    for e in K.ENTITY_ORDER:
        line = K.ENTITIES[e]["line"]
        for ino in a[e]["missing"]:
            rs = [x for x in rows if x["incident_no"] == ino and x["coverage_line"] == line]
            opn = sorted((x for x in rs if x["status"] in ("OPEN", "REOPENED")),
                         key=lambda x: (x["claim_made_date"], x["claim_no"]))[0]
            res = sum((x["indemnity_reserve"] or 0) + (x["defense_reserve"] or 0) for x in rs)
            for col, val in enumerate([K.ENTITIES[e]["name"], opn["claim_no"], ino,
                                       opn["incident_date"].strftime("%d %b %Y").lstrip("0"),
                                       opn["claim_made_date"].strftime("%d %b %Y").lstrip("0"),
                                       opn["matter_type"], r1000(res)], start=1):
                cell = ws3.cell(row=rr, column=col, value=val)
                if col == 7:
                    cell.number_format = "#,##0"
            rr += 1
            n_missing += 1
    if not n_missing:
        ws3.cell(row=rr, column=1, value="None")
        rr += 1
    rr += 1
    for t in ("Each expiring carrier's policies are the terms it wrote for the entity, from the date in the premium "
              "history; open claims from incidents before that date are under the former carrier's occurrence policies "
              "(trust agreement 2.2), so they are not looked for on these runs.",
              "Quillan Mutual's latest run is the one valued 30 September 2026 (issued 6 October 2026); revision 1 of "
              "its 30 June 2026 run, issued 9 October 2026, is an earlier valuation.",
              "Claims that began as precautionary notices are carried on the runs under their notice numbers (RMN), "
              "which the notice log ties to incident numbers; a suit that followed an earlier claim number for the "
              "same incident stays under that earlier number. Stoneleigh drops the punctuation from references.",
              "Missing: open in the register at 30 September 2026 on the entity's coverage line, one per incident, "
              "with no record on the run. Reserves are the register's indemnity and defense reserves on our rows."):
        ws3.cell(row=rr, column=1, value=t).font = NOTE_FONT
        rr += 1
    for col, wd in zip(range(1, 8), [30, 44, 22, 24, 18, 18, 22]):
        ws3.column_dimensions[get_column_letter(col)].width = wd

    ws4 = wb.create_sheet("Covered FTE 2026")
    ws4["A1"] = "Covered clinician FTE paid by the payroll, July to September 2026"
    ws4["A1"].font = T
    _head(ws4, 3, ["Participating Entity"] + [MONTH_LABEL[x] for x in A.ASK_B_MONTHS])
    for i, e in enumerate(K.ENTITY_ORDER):
        ws4.cell(row=4 + i, column=1, value=K.ENTITIES[e]["name"])
        for j, mo in enumerate(A.ASK_B_MONTHS):
            cell = ws4.cell(row=4 + i, column=2 + j, value=r1(b["gold"][(e, mo)]))
            cell.number_format = "#,##0.0"
    ws4.cell(row=8, column=1, value="All entities").font = B
    for j, mo in enumerate(A.ASK_B_MONTHS):
        cell = ws4.cell(row=8, column=2 + j, value=r1(sum(b["gold"][(e, mo)] for e in K.ENTITY_ORDER)))
        cell.number_format = "#,##0.0"
        cell.font = B
        cell.fill = TOTAL
    ws4.cell(row=8, column=1).fill = TOTAL
    for k, t in enumerate(("Paid FTE: the hours paid on each covered position's pay lines in the work month (regular "
                           "and paid time off) over the 173.33 hours scheduled a month at 1.0 FTE; a per-diem "
                           "clinician by the shifts paid, a month of unpaid leave at none.",
                           "Covered Clinicians: Schedule B jobs through the HCM job profile crosswalk. A resident's or "
                           "fellow's additional position is left out, trainees not being Covered Clinicians (trust "
                           "agreement 1.5).",
                           "Entity: the entity of the cost centre the lines are charged to. The total is rounded from "
                           "unrounded figures.",
                           "Source: HCM clinician pay lines for work months January to September 2026, run 2 October "
                           "2026.")):
        ws4.cell(row=10 + k, column=1, value=t).font = NOTE_FONT
    for col, wd in zip(range(1, 5), [34, 16, 16, 18]):
        ws4.column_dimensions[get_column_letter(col)].width = wd

    ws5 = wb.create_sheet("Notes")
    ws5["A1"] = "Notes"
    ws5["A1"].font = T
    notes = [
        ("Claim", "As the Trust Agreement defines it (1.3): every demand, Early Resolution Program request or suit "
                  "from one Medical Incident is one Claim, first made on the date the earliest of them reached Risk "
                  "Management."),
        ("Covered Clinician", "Employed physicians and advanced practice clinicians in a Schedule B job. Residents, "
                              "fellows, psychologists and contracted or locum clinicians are excluded."),
        ("Window", "Each entity from its Transition Date (its Retroactive Date) to 31 December 2027: Claims made in "
                   "2027 from incidents on or after that date."),
        ("Experience", "The %s Claims from incidents in 2016 to 2024 naming Covered Clinicians, valued at paid plus "
                       "case reserves on the Covered Clinicians' rows of the register extract valued 30 September "
                       "2026 and trended at 5.0%% a year from the incident year to 2027, against %s covered FTE-years "
                       "in the payroll for the same years, pooled within each specialty class (funding policy 3.2 to "
                       "3.4)." % ("{:,}".format(m["exp_claims"]), "{:,.1f}".format(r1(m["exp_fte_years"])))),
        ("Exposure", "The staffing plan's covered FTE for each month from the Transition Date, a month one twelfth of "
                     "an FTE-year, held through the month's days."),
        ("Expected cost", "Each experience Claim stands for the Claims its class would make in 2027 at the same delay "
                          "from incident to first making: it counts once for each FTE-year of 2027 exposure with "
                          "that delay still left before 31 December, per FTE-year of its class's experience, at its "
                          "own cost."),
        ("Not included", "Administrative expenses, investment income, excess insurance and any margin (funding policy "
                         "3.5)."),
        ("Rounding", "Rounded once, from unrounded amounts: the contribution to the nearest $250,000, other dollars to "
                     "the nearest thousand, FTE to one decimal."),
        ("As of", "Register extract run 2 October 2026, valued 30 September 2026; HCM pay lines to September 2026; "
                  "carrier loss runs received 5 to 9 October 2026."),
    ]
    for i, (k, t) in enumerate(notes, start=3):
        ws5.cell(row=i, column=1, value=k).font = B
        cell = ws5.cell(row=i, column=2, value=t)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws5.cell(row=i, column=1).alignment = Alignment(vertical="top")
        ws5.row_dimensions[i].height = 15 * (1 + len(t) // 95)
    ws5.column_dimensions["A"].width = 18
    ws5.column_dimensions["B"].width = 100
    for sheet in wb.worksheets:
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)
    PP.normalize_ooxml(path, K.PEOPLE["requester"], K.PEOPLE["requester"], "2026-10-12T17:05:31Z",
                       "2026-10-14T00:41:09Z", excel=True)


# ------------------------------------------------------------------------------------------------ deck
def chart_title(b):
    tot = {mo: sum(b["gold"][(e, mo)] for e in K.ENTITY_ORDER) for mo in A.ASK_B_MONTHS}
    big = max(K.ENTITY_ORDER, key=lambda e: b["gold"][(e, "2026-09")])
    share = b["gold"][(big, "2026-09")] / tot["2026-09"]
    return "%s carries %d%% of the covered clinician FTE the payroll paid in September" % (
        K.ENTITIES[big]["name"], int(math.floor(share * 100 + 0.5)))


def write_chart(buf, c):
    b = c["b"]
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "svg.hashsalt": "t130",
                         "axes.edgecolor": SLATE, "axes.labelcolor": SLATE, "xtick.color": NAVY, "ytick.color": SLATE})
    fig, ax = plt.subplots(figsize=(9.6, 4.6), dpi=200)
    fig.patch.set_facecolor(PAPER)
    ax.set_facecolor(PAPER)
    width = 0.26
    xs = list(range(len(K.ENTITY_ORDER)))
    top = max(b["gold"].values())
    for j, mo in enumerate(A.ASK_B_MONTHS):
        vals = [r1(b["gold"][(e, mo)]) for e in K.ENTITY_ORDER]
        pos = [x + (j - 1) * width for x in xs]
        ax.bar(pos, vals, width=width * 0.94, color=MONTH_COLOURS[mo], label=MONTH_LABEL[mo], edgecolor=PAPER,
               linewidth=0.6)
        for p, v in zip(pos, vals):
            ax.text(p, v + top * 0.012, "{:,.1f}".format(v), ha="center", va="bottom", fontsize=7.6, color=NAVY)
    ax.set_xticks(xs)
    ax.set_xticklabels([K.ENTITIES[e]["name"] for e in K.ENTITY_ORDER])
    ax.set_ylabel("Covered clinician FTE paid")
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: "{:,.0f}".format(v)))
    ax.yaxis.grid(True, color=LINE, linewidth=0.8)
    ax.set_axisbelow(True)
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    ax.tick_params(axis="x", length=0)
    ax.set_ylim(0, top * 1.12)
    ax.legend(loc="upper left", frameon=False, fontsize=8.8)
    fig.tight_layout()
    fig.savefig(buf, format="png", facecolor=PAPER, metadata={"Software": None})
    plt.close(fig)


def write_deck(path, c, chart_png):
    m, b = c["m"], c["b"]
    navy, slate, pale = RGBColor(0x1F, 0x2F, 0x3F), RGBColor(0x5A, 0x66, 0x72), RGBColor(0xEE, 0xF1, 0xF4)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
    blank = prs.slide_layouts[6]

    def text(slide, x, y, w, h, runs, size=12, color=slate, bold=False):
        tf = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h)).text_frame
        tf.word_wrap = True
        for k, line in enumerate(runs if isinstance(runs, list) else [runs]):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = line
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
        return tf

    def frame(slide, n, title, sub):
        text(slide, 1.6, 0.9, 30.6, 1.6, title, size=24, color=navy, bold=True)
        text(slide, 1.6, 2.5, 30.6, 1.0, sub, size=13)
        text(slide, 1.6, 17.9, 30.6, 0.7, "Corvane Health, Risk Management and Insurance  |  Trustees of the Corvane "
             "Professional Liability Trust, 10 November 2026  |  %d" % n, size=9)

    contribution = "{:,}".format(m["rounded"])
    s1 = prs.slides.add_slide(blank)
    frame(s1, 1, "2027 contribution to the Trust: $%s" % contribution,
          "Recommended: the Participating Entities contribute $%s for 2027, the expected cost of the Claims that will "
          "be first made against the Trust in its first Plan Year" % contribution)
    rows = [["Participating Entity", "Joins the Trust", "Months in the Trust in 2027"]]
    for e in K.ENTITY_ORDER:
        rows.append([K.ENTITIES[e]["name"], long_date(K.ENTITIES[e]["transition"]), "%d" % months_in(e)])
    widths = [9.4, 6.4, 7.6]
    tb = s1.shapes.add_table(len(rows), 3, Cm(1.6), Cm(4.3), Cm(sum(widths)), Cm(0.95 * len(rows))).table
    for j, wd in enumerate(widths):
        tb.columns[j].width = Cm(wd)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tb.cell(i, j)
            cell.text = val
            cell.fill.solid()
            last = False
            cell.fill.fore_color.rgb = navy if i == 0 else (pale if last else RGBColor(0xFF, 0xFF, 0xFF))
            para = cell.text_frame.paragraphs[0]
            if not para.runs:
                para.add_run()
            run = para.runs[0]
            run.font.size = Pt(12 if i else 11)
            run.font.bold = i == 0 or last
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else navy
            if j >= 2:
                para.alignment = PP_ALIGN.RIGHT
    text(s1, 1.6, 10.6, 30.6, 6.6, [
        "The Trust is claims-made, each entity's Transition Date its Retroactive Date: in 2027 it pays only for the "
        "Claims made by 31 December from incidents after the entity joined.",
        "Those are the Claims the System's experience makes within the year: harm disclosed and settled within weeks "
        "of the care, which is costly, and program requests a few months on, which settle small. The demands and "
        "suits that carry most of the experience's cost come a year or more after the care, after 2027 for every "
        "entity.",
        "Each 2016 to 2024 experience Claim was costed at its own delay from incident to first making against the "
        "days each entity is in the Trust in 2027, within its specialty class.",
        "No provision for expenses, investment income, excess insurance or margin (funding policy 3.5)."], size=12)

    s2 = prs.slides.add_slide(blank)
    frame(s2, 2, chart_title(b),
          "Covered clinician FTE the payroll paid at each entity, July to September 2026 (HCM pay lines)")
    s2.shapes.add_picture(chart_png, Cm(2.6), Cm(3.5), width=Cm(27.6))
    text(s2, 1.6, 16.9, 30.6, 0.8, "Source: HCM clinician pay lines run 2 October 2026; hours paid over 173.33 a month "
         "at 1.0 FTE, residents' and fellows' additional positions left out, entity of the cost centre charged.",
         size=9)
    prs.core_properties.title = "Trust contribution 2027"
    prs.save(path)
    PP.normalize_ooxml(path, K.PEOPLE["requester"], K.PEOPLE["requester"], "2026-10-13T16:22:05Z",
                       "2026-10-14T01:10:47Z", title="Trust contribution 2027")


def main(target, golden, figures_path, world=None):
    os.makedirs(golden, exist_ok=True)
    c = compute(target)
    fig = figures(c)
    write_book(os.path.join(golden, BOOK), c)
    buf = io.BytesIO()
    write_chart(buf, c)
    buf.seek(0)
    write_deck(os.path.join(golden, DECK), c, buf)
    with open(figures_path, "w") as fh:
        json.dump(fig, fh, indent=1, ensure_ascii=False, default=str)
        fh.write("\n")
    return c


if __name__ == "__main__":
    task = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    main(os.path.join(task, "target"), os.path.join(task, "golden"), os.path.join(task, "generator", "figures.json"))
