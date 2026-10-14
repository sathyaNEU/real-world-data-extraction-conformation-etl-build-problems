"""The pack's documents: the trust agreement (executed extract), the trustees' funding policy, the 2027 clinician
staffing plan, Quillan Mutual's two loss runs, the premium history and the data dictionary for the risk management
extracts."""
from __future__ import annotations

import datetime as dt

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import LETTER                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import inch                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import common as K   # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for _n, _f in (("LiberationSerif", "LiberationSerif-Regular.ttf"), ("LiberationSerif-Bold", "LiberationSerif-Bold.ttf"),
               ("LiberationSerif-Italic", "LiberationSerif-Italic.ttf"),
               ("LiberationSans", "LiberationSans-Regular.ttf"), ("LiberationSans-Bold", "LiberationSans-Bold.ttf")):
    pdfmetrics.registerFont(TTFont(_n, FONT_DIR + _f))
SERIF, SERIF_B, SERIF_I, SANS, SANS_B = ("LiberationSerif", "LiberationSerif-Bold", "LiberationSerif-Italic",
                                         "LiberationSans", "LiberationSans-Bold")
INK = colors.HexColor("#1b2127")
GREY = colors.HexColor("#5b636b")
RULE = colors.HexColor("#97a1ab")
BODY = ParagraphStyle("body", fontName=SERIF, fontSize=10.4, leading=13.9, textColor=INK, spaceAfter=5)
CLAUSE = ParagraphStyle("clause", parent=BODY, leftIndent=30, firstLineIndent=-30)
SUB = ParagraphStyle("sub", parent=BODY, leftIndent=48, firstLineIndent=-18)
TITLE = ParagraphStyle("title", fontName=SERIF_B, fontSize=15, leading=19, textColor=INK, alignment=1, spaceAfter=3)
SUBTITLE = ParagraphStyle("subtitle", fontName=SERIF, fontSize=10.5, leading=13.5, textColor=INK, alignment=1,
                          spaceAfter=12)
H2 = ParagraphStyle("h2", fontName=SERIF_B, fontSize=11.2, leading=14.5, textColor=INK, spaceBefore=9, spaceAfter=4)
NOTE = ParagraphStyle("note", fontName=SERIF_I, fontSize=9.2, leading=11.8, textColor=GREY, spaceBefore=2,
                      spaceAfter=6)
CELL = ParagraphStyle("cell", fontName=SANS, fontSize=8.3, leading=10.1, textColor=INK)
CELLB = ParagraphStyle("cellb", parent=CELL, fontName=SANS_B)


def _frame(left, right):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont(SANS, 7.6)
        canvas.setFillColor(GREY)
        canvas.drawString(1.0 * inch, 0.58 * inch, left)
        canvas.drawRightString(LETTER[0] - 1.0 * inch, 0.58 * inch, "%s  |  Page %d" % (right, doc.page))
        canvas.restoreState()
    return draw


def _doc(path, title, author):
    return SimpleDocTemplate(path, pagesize=LETTER, leftMargin=1.0 * inch, rightMargin=1.0 * inch,
                             topMargin=0.85 * inch, bottomMargin=0.9 * inch, title=title, author=author,
                             subject="", creator="", producer="")


def _table(rows, widths, head=1):
    data = [[Paragraph(str(c), CELLB if i < head else CELL) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, repeatRows=head)
    t.setStyle(TableStyle([("LINEBELOW", (0, head - 1), (-1, head - 1), 0.6, RULE),
                           ("LINEBELOW", (0, -1), (-1, -1), 0.4, RULE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("TOPPADDING", (0, 0), (-1, -1), 2.0), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.0),
                           ("LEFTPADDING", (0, 0), (-1, -1), 3)]))
    return t


def long_date(d):
    return "%d %s %d" % (d.day, d.strftime("%B"), d.year)


def expiring_policy(e):
    t = K.ENTITIES[e]["transition"]
    last = t - dt.timedelta(days=1)
    return K.policy_number(e, last), K.policy_period(e, last)


# ------------------------------------------------------------------------------------------------ trust agreement
def write_agreement(path):
    doc = _doc(path, "Corvane Professional Liability Trust Agreement", "Corvane Health")
    s = [Paragraph("CORVANE PROFESSIONAL LIABILITY TRUST", TITLE),
         Paragraph("Trust Agreement between Corvane Health and the Trustees, executed 14 August 2026<br/>"
                   "Extract: Articles 1, 2 and 5 and Schedules A and B", SUBTITLE)]
    s.append(Paragraph("Article 1. Definitions", H2))
    for no, text in [
        ("1.1", "<b>System</b> means Corvane Health, an Oregon nonprofit corporation, and the Participating Entities."),
        ("1.2", "<b>Participating Entity</b> means each entity listed in Schedule A."),
        ("1.3", "<b>Claim</b> means (a) a written demand against a Covered Clinician for money or services arising "
                "from a Medical Incident, (b) a request for compensation made through the System's Early Resolution "
                "Program arising from a Medical Incident, or (c) the service of a civil proceeding against a Covered "
                "Clinician arising from a Medical Incident. All Claims arising from the same Medical Incident are one "
                "Claim, first made on the date the earliest of them was made. A Claim is made on the date it is first "
                "received by the System's Risk Management and Insurance department."),
        ("1.4", "<b>Medical Incident</b> means an act or omission in the rendering of, or failure to render, "
                "professional health care services."),
        ("1.5", "<b>Covered Clinician</b> means a physician or advanced practice clinician employed by a Participating "
                "Entity in a job listed in Schedule B, while so employed. Residents, fellows and other trainees, and "
                "clinicians engaged by contract or as locum tenens, are not Covered Clinicians."),
        ("1.6", "<b>Transition Date</b> and <b>Retroactive Date</b> mean, for each Participating Entity, the date "
                "shown for it in Schedule A."),
        ("1.7", "<b>Plan Year</b> means a calendar year. The first Plan Year is 2027."),
        ("1.8", "<b>Trustees</b> means the board of trustees of the Trust appointed under Article 3."),
    ]:
        s.append(Paragraph("%s&nbsp;&nbsp;&nbsp;%s" % (no, text), CLAUSE))
    s.append(Paragraph("Article 2. Coverage", H2))
    for no, text in [
        ("2.1", "The Trust will pay on behalf of a Covered Clinician the damages and defence costs the Covered "
                "Clinician becomes legally obligated to pay by reason of a Claim first made against the Covered "
                "Clinician during a Plan Year, arising from a Medical Incident that occurred on or after the "
                "Retroactive Date of the Participating Entity that employed the Covered Clinician when the Medical "
                "Incident occurred."),
        ("2.2", "A Medical Incident that occurred before a Participating Entity's Transition Date remains insured "
                "under the occurrence policy in force on the date it occurred. Schedule A lists each Participating "
                "Entity's expiring policy. The Trust provides no prior acts coverage."),
        ("2.3", "The Trust's limits are $2,000,000 each Claim and $4,000,000 in the aggregate for each Covered "
                "Clinician in each Plan Year, defence costs included."),
        ("2.4", "A Claim first made after the end of a Plan Year is not covered by that Plan Year. The Trust provides "
                "an extended reporting period only on termination under Article 9."),
    ]:
        s.append(Paragraph("%s&nbsp;&nbsp;&nbsp;%s" % (no, text), CLAUSE))
    s.append(Paragraph("Article 5. Claims administration", H2))
    for no, text in [
        ("5.1", "The System's Risk Management and Insurance department administers Claims for the Trust, keeps the "
                "claims register and reports to the Trustees each quarter."),
        ("5.2", "The System's Early Resolution Program is open to the patients of every Participating Entity from its "
                "Transition Date."),
        ("5.3", "Each Participating Entity continues to report to its expiring carrier every Claim and every "
                "precautionary notice of a Medical Incident that occurred before its Transition Date, for as long as "
                "that carrier's policy responds."),
        ("5.4", "Defence counsel is appointed from the panel the Trustees approve. No settlement above $250,000 is made "
                "without the consent of the Trustees' claims committee."),
    ]:
        s.append(Paragraph("%s&nbsp;&nbsp;&nbsp;%s" % (no, text), CLAUSE))
    s.append(Paragraph("Schedule A. Participating Entities", H2))
    rows = [["Participating Entity", "Code", "Transition Date and Retroactive Date", "Expiring occurrence policy",
             "Policy period"]]
    for e in sorted(K.ENTITY_ORDER, key=lambda x: K.ENTITIES[x]["transition"]):
        ent = K.ENTITIES[e]
        pol, (a, b) = expiring_policy(e)
        rows.append([ent["name"], e, long_date(ent["transition"]), "%s, %s" % (K.CARRIERS[ent["carrier"]], pol),
                     "%s to %s" % (long_date(a), long_date(b))])
    s.append(_table(rows, [1.35 * inch, 0.45 * inch, 1.2 * inch, 1.85 * inch, 1.65 * inch]))
    s.append(Paragraph("Each expiring policy is the last term of the entity's annual occurrence programme with its "
                       "carrier and is not renewed. A policy period ends at 12:01 a.m. on its last date.", NOTE))
    s.append(Paragraph("Schedule B. Covered jobs and specialty classes", H2))
    rows = [["Specialty class", "Job code to June 2016", "Job code from July 2016", "Job title"]]
    for cls in K.CLASSES:
        for j in K.JOBS:
            if j[4] == cls:
                rows.append([K.CLASS_NAME[cls], j[1], j[2], j[3]])
    s.append(_table(rows, [1.85 * inch, 1.15 * inch, 1.2 * inch, 2.3 * inch]))
    s.append(Paragraph("Note 1. Jobs not listed are not covered, among them resident physicians (25010, GME-RES), "
                       "clinical fellows (25020, GME-FEL) and psychologists (32110, BHV-PSY).", NOTE))
    s.append(Paragraph("Note 2. A Claim takes the specialty class of the first Covered Clinician named in it, by that "
                       "clinician's job in the month of the Medical Incident.", NOTE))
    s.append(Spacer(1, 8))
    s.append(Paragraph("Executed for Corvane Health by %s, Chief Financial Officer, and for the Trustees by %s, Chair, "
                       "on 14 August 2026." % (K.PEOPLE["cfo"], K.PEOPLE["chair"]), BODY))
    frame = _frame("Corvane Professional Liability Trust Agreement (extract)", "Executed 14 August 2026")
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ------------------------------------------------------------------------------------------------ funding policy
def write_policy(path):
    doc = _doc(path, "Corvane Professional Liability Trust Funding Policy", "Corvane Professional Liability Trust")
    s = [Paragraph("CORVANE PROFESSIONAL LIABILITY TRUST", TITLE),
         Paragraph("Funding Policy, adopted by the Trustees on 22 September 2026", SUBTITLE)]
    s.append(Paragraph("1. Purpose", H2))
    s.append(Paragraph("This policy sets how the Trustees fix the contributions the Participating Entities make to the "
                       "Trust for each Plan Year. Capitalised terms have the meanings given in the Trust Agreement.",
                       BODY))
    s.append(Paragraph("2. Timing", H2))
    s.append(Paragraph("The Trustees fix the Contribution for a Plan Year at their November meeting in the year before "
                       "it, on the recommendation of the System's Vice President, Risk Management and Insurance. The "
                       "Participating Entities pay it in equal quarterly instalments, each in proportion to its share "
                       "of the Plan Year's exposure under 3.3.", BODY))
    s.append(Paragraph("3. The Contribution", H2))
    for no, text in [
        ("3.1", "For each Plan Year the Participating Entities together contribute the expected cost, at the Plan "
                "Year's cost level, of the Claims that will be first made against the Trust during the Plan Year "
                "under the Trust Agreement (the Contribution)."),
        ("3.2", "The expected cost is estimated from the System's own claims experience: the Claims arising from "
                "Medical Incidents from 1 January 2017 to 31 December 2024 against Covered Clinicians, each valued at "
                "the indemnity and defence paid and the case reserves recorded against its Covered Clinicians in the "
                "risk management claims register extract valued 30 September 2026, and brought to the Plan Year's "
                "cost level at 5.0 per cent a year compounded from the calendar year of the Medical Incident. The "
                "exposure behind that experience is the Covered Clinicians' paid FTE in the System's payroll for the "
                "same years, each month's FTE counting as one twelfth of an FTE-year."),
        ("3.3", "The Plan Year's exposure is the Covered Clinician FTE in the System's staffing plan for the Plan "
                "Year, for each Participating Entity, specialty class and month from the entity's Transition Date. "
                "Each month's FTE counts as one twelfth of an FTE-year, earned evenly over the days of the month."),
        ("3.4", "The Trust rates its Participating Entities alike: a Covered Clinician's exposure in a specialty "
                "class carries the same expected cost whichever Participating Entity employs the clinician."),
        ("3.5", "Each Claim in the experience is taken at its value as recorded in the extract: no development, tail "
                "or credibility factor is applied and no distribution is fitted."),
        ("3.6", "The first Plan Year's Contribution carries no provision for the Trust's administrative expenses, "
                "investment income, excess insurance or a margin; the Trustees budget those separately."),
        ("3.7", "Figures reported to the Trustees are rounded once, from unrounded amounts; a total is the rounded sum "
                "of its unrounded parts."),
    ]:
        s.append(Paragraph("%s&nbsp;&nbsp;&nbsp;%s" % (no, text), CLAUSE))
    s.append(Paragraph("4. Review", H2))
    s.append(Paragraph("After each Plan Year closes, the Trust's consulting actuary reviews the Claims made in it "
                       "against the Contribution and reports to the Trustees by 30 June. The Trustees review this "
                       "policy every two years, and when a Participating Entity joins or leaves the Trust.", BODY))
    s.append(Paragraph("5. Investment of contributions", H2))
    s.append(Paragraph("Contributions are held in the Trust's operating account and its short-term investment pool "
                       "under the System's investment policy for operating funds. No contribution is invested for a "
                       "term longer than two years.", BODY))
    s.append(Spacer(1, 10))
    s.append(Paragraph("Adopted by the Trustees on 22 September 2026. %s, Chair." % K.PEOPLE["chair"], BODY))
    frame = _frame("Corvane Professional Liability Trust Funding Policy", "Adopted 22 September 2026")
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ------------------------------------------------------------------------------------------------ workbooks
HEAD_FILL = PatternFill("solid", fgColor="D9E1EA")
BOLD = Font(name="Calibri", size=11, bold=True)
TITLE_FONT = Font(name="Calibri", size=13, bold=True)
MONTHS = ["Jan-27", "Feb-27", "Mar-27", "Apr-27", "May-27", "Jun-27", "Jul-27", "Aug-27", "Sep-27", "Oct-27", "Nov-27",
          "Dec-27"]


def write_plan(path, plan):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Covered FTE 2027"
    ws["A1"] = "Corvane Health, clinician staffing plan 2027"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Covered Clinician FTE by Participating Entity and specialty class (Trust Agreement Schedule B), budgeted " \
               "FTE for each month"
    ws["A3"] = "Financial Planning and Analysis, plan version 2026-09-29"
    hdr = ["Entity", "Entity name", "Specialty class"] + MONTHS
    for c, h in enumerate(hdr, start=1):
        cell = ws.cell(row=5, column=c, value=h)
        cell.font = BOLD
        cell.fill = HEAD_FILL
    r = 6
    for e in K.ENTITY_ORDER:
        for cls in K.CLASSES:
            if (e, cls) not in plan:
                continue
            ws.cell(row=r, column=1, value=e)
            ws.cell(row=r, column=2, value=K.ENTITIES[e]["name"])
            ws.cell(row=r, column=3, value=K.CLASS_NAME[cls])
            for k, v in enumerate(plan[(e, cls)]):
                cell = ws.cell(row=r, column=4 + k, value=float(v))
                cell.number_format = "0.0"
            r += 1
    for c, w in zip(range(1, 16), [8, 28, 32] + [8.5] * 12):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "D6"
    ns = wb.create_sheet("Notes")
    notes = ["Plan FTE are budgeted positions at each month's planned level; a vacancy is carried from the month the "
             "plan expects it filled.",
             "Covered Clinician FTE only: resident physicians, clinical fellows, psychologists and contracted or locum "
             "clinicians are not included.",
             "Entity codes as in the general ledger: TFH Tolliver Falls Hospital, AMC Alderwick Medical Center, CMG "
             "Corvane Medical Group, BCH Brennick Community Hospital.",
             "Approved with the 2027 operating budget by the Finance Committee on 29 September 2026."]
    ns["A1"] = "Notes"
    ns["A1"].font = BOLD
    for k, t in enumerate(notes, start=2):
        ns.cell(row=k, column=1, value=t)
    ns.column_dimensions["A"].width = 120
    wb.save(path)


QM_HEADER = ["Policy Number", "Policy Period", "Named Insured", "Claim Number", "Insured Reference", "Date of Loss",
             "Date Reported", "Claim Made", "Claim Type", "Status", "Indemnity Paid", "Expense Paid",
             "Indemnity Reserve", "Expense Reserve", "Total Incurred"]


def write_qm_run(path, recs, nums, valued, run_date, revision=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Loss Run"
    ws["A1"] = K.CARRIERS["QM"]
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Hospital Professional Liability, Loss Run Report"
    ws["A3"] = "Account: Corvane Health (Tolliver Falls Hospital; Alderwick Medical Center)"
    ws["A4"] = "Valuation Date: %s" % valued.strftime("%m/%d/%Y")
    ws["A5"] = "Run Date: %s%s" % (run_date.strftime("%m/%d/%Y"), ("    " + revision) if revision else "")
    for c, h in enumerate(QM_HEADER, start=1):
        cell = ws.cell(row=7, column=c, value=h)
        cell.font = BOLD
        cell.fill = HEAD_FILL
    r = 8
    names = {"TFH": "Tolliver Falls Hospital", "AMC": "Alderwick Medical Center"}
    for e in ("AMC", "TFH"):
        block = sorted((x for x in recs if x["entity"] == e), key=lambda x: (x["policy"], nums[x["incident_no"]]))
        tot = [0.0] * 5
        for x in block:
            a, b = x["period"]
            vals = [x["policy"], "%s-%s" % (a.strftime("%m/%d/%Y"), b.strftime("%m/%d/%Y")), names[e],
                    nums[x["incident_no"]], x["ref"], x["date_of_loss"], x["created"], x["claim_rptd"], x["ctype"],
                    x["status"], x["paid_ind"], x["paid_exp"], x["res_ind"], x["res_exp"],
                    round(x["paid_ind"] + x["paid_exp"] + x["res_ind"] + x["res_exp"], 2)]
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=c, value=v)
                if c in (6, 7, 8):
                    cell.number_format = "mm/dd/yyyy"
                if c >= 11:
                    cell.number_format = "#,##0.00"
            for k in range(5):
                tot[k] += vals[10 + k]
            r += 1
        ws.cell(row=r, column=3, value="Total %s (%d claims and incidents)" % (names[e], len(block))).font = BOLD
        for k in range(5):
            cell = ws.cell(row=r, column=11 + k, value=round(tot[k], 2))
            cell.number_format = "#,##0.00"
            cell.font = BOLD
        r += 2
    for c, w in zip(range(1, 16), [17, 23, 24, 13, 17, 12, 13, 12, 11, 10, 15, 14, 16, 15, 15]):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A8"
    wb.save(path)


def write_premiums(path):
    rng = K.stream("premiums")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Premium history"
    ws["A1"] = "Corvane Health, clinician professional liability premium history 2017 to 2026"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Risk Management and Insurance; annual premiums as billed at inception, before audit adjustments"
    hdr = ["Entity", "Carrier", "Policy number", "Policy period start", "Policy period end", "Limits",
           "Rated physician FTE", "Rated APC FTE", "Annual premium ($)"]
    for c, h in enumerate(hdr, start=1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = BOLD
        cell.fill = HEAD_FILL
    base = {"TFH": (52.0, 27.0, 1.62e6), "AMC": (148.0, 58.0, 4.9e6), "CMG": (338.0, 151.0, 5.6e6),
            "BCH": (26.0, 19.0, 0.62e6)}
    r = 5
    for e in K.ENTITY_ORDER:
        ent = K.ENTITIES[e]
        for y in range(2017, 2027):
            start = dt.date(y, ent["anniv"], 1)
            if ent["joined"] and start < ent["joined"]:
                continue
            phys, apc, prem = base[e]
            g = (1.022 ** (y - 2026))
            rated_p = round(phys * g * float(rng.uniform(0.97, 1.03)), 1)
            rated_a = round(apc * g * float(rng.uniform(0.95, 1.05)), 1)
            premium = round(prem * (1.045 ** (y - 2026)) * float(rng.uniform(0.96, 1.05)), -2)
            vals = [ent["name"], K.CARRIERS[ent["carrier"]], K.policy_number(e, start), start,
                    dt.date(y + 1, ent["anniv"], 1), "$1,000,000 / $3,000,000", rated_p, rated_a, premium]
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=c, value=v)
                if c in (4, 5):
                    cell.number_format = "yyyy-mm-dd"
                if c == 9:
                    cell.number_format = "#,##0"
            r += 1
    for c, w in zip(range(1, 10), [30, 44, 20, 18, 18, 22, 18, 14, 18]):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A5"
    wb.save(path)


# ------------------------------------------------------------------------------------------------ data dictionary
def write_dictionary(path, counts):
    t = []
    a = t.append
    a("# Risk management extracts: data dictionary")
    a("")
    a("Corvane Health, Risk Management and Insurance. Field definitions for the files extracted on 2 October 2026 for "
      "the 2027 trust contribution. Maintained by %s's claims team; questions to the RMI analytics mailbox." %
      K.PEOPLE["claims"])
    a("")
    a("## %s" % K.REGISTER)
    a("")
    a("One row per clinician named as a defendant on a claim number. The register holds every claim naming a clinician "
      "who practised at a Corvane entity, whoever insures the clinician; claims against the entities alone are kept "
      "in the general liability register. Valued at 30 September 2026; %d rows." % counts[K.REGISTER])
    a("")
    for f, d in [
        ("claim_no", "Risk management claim number, CL-year-sequence, given when the claim is received."),
        ("incident_no", "Risk management incident number of the event the claim arises from, IR-year-sequence."),
        ("defendant_seq", "Order in which the clinician is named on the claim number."),
        ("clinician_id", "Payroll employee id for employed clinicians, residents and fellows; CTR- credentialing number "
                         "for contracted and locum clinicians."),
        ("defendant_role", "Attending physician, Advanced practice clinician, Resident physician, Clinical fellow or "
                           "Contracted physician."),
        ("specialty", "Specialty as entered by the claims analyst."),
        ("coverage_line", "The insurance programme responding for this defendant: carrier code and entity (QM-TFH, "
                          "QM-AMC, SC-CMG, SC-BCH); CONTRACT where the clinician's own group insures them."),
        ("entity_at_incident", "Entity where the care was given."),
        ("incident_date", "Date of the care the claim arises from."),
        ("claim_made_date", "Date risk management received the claim under this claim number."),
        ("matter_type", "Claim (a demand for compensation, or a request through the Early Resolution Program, which "
                        "risk management has run for the patients of Tolliver Falls Hospital and Alderwick Medical "
                        "Center since September 2015) or Suit (a civil proceeding served)."),
        ("allegation", "Principal allegation."),
        ("status", "OPEN, CLOSED or REOPENED at the valuation date. A reopened claim keeps its earlier closing date."),
        ("closed_date", "Date the claim number was last closed."),
        ("indemnity_paid, defense_paid", "Payments to date on this defendant, in dollars. Blank on CONTRACT rows."),
        ("indemnity_reserve, defense_reserve", "Case reserves on this defendant at the valuation date. Blank on "
                                               "CONTRACT rows."),
        ("prior_claim_no", "Where the claim number was opened from an earlier claim number for the same matter, the "
                           "earlier claim number."),
        ("valuation_date", "Valuation date of the extract."),
    ]:
        a("- `%s`: %s" % (f, d))
    a("")
    a("## %s" % K.PAYROLL)
    a("")
    a("Legacy payroll history for employed clinicians, residents, fellows and psychologists, January 2016 to December "
      "2025 (the legacy payroll closed with the December 2025 pay run). One row per employee, month and cost centre; "
      "%d rows." % counts[K.PAYROLL])
    a("")
    for f, d in [
        ("pay_month", "Month paid, YYYY-MM."),
        ("employee_id", "Payroll employee id."),
        ("job_code, job_title", "Job held in the month. Codes changed with the job architecture project on 1 July 2016."),
        ("cost_centre", "Company and department, ccc-dddd."),
        ("entity_code", "Entity of the cost centre."),
        ("fte", "Paid FTE for the month on this cost centre."),
        ("emp_status", "A active, L leave of absence, T terminated in the month."),
    ]:
        a("- `%s`: %s" % (f, d))
    a("")
    a("## %s" % K.HCM)
    a("")
    a("Pay lines from the HCM system that replaced the legacy payroll on 1 January 2026, work months January to "
      "September 2026. One row per worker, position, work month and earning code; %d rows." % counts[K.HCM])
    a("")
    for f, d in [
        ("worker_id", "Employee id (carried over from the legacy payroll)."),
        ("position_id", "HCM position."),
        ("work_month", "Month worked, YYYY-MM."),
        ("earning_code, earning_description", "REG, PTO, CALL, ADMIN or LOA."),
        ("hours", "Hours paid on the line."),
        ("fte", "FTE of the position."),
        ("job_profile_code, job_profile", "HCM job profile of the position; legacy job codes through the crosswalk."),
        ("cost_centre", "Cost centre the line is charged to."),
        ("home_entity", "The worker's home entity on their HCM record."),
        ("pay_group", "PHYS, APC or GME."),
    ]:
        a("- `%s`: %s" % (f, d))
    a("")
    a("## %s and %s" % (K.CROSSWALK, K.DIRECTORY))
    a("")
    a("The HCM job profile crosswalk maps each HCM job profile to the legacy job code (July 2016 code set) it replaced. The "
      "cost centre directory gives each cost centre's company, entity and department with the dates it is effective.")
    a("")
    a("## %s" % K.NOTICES)
    a("")
    a("One row per precautionary notice risk management sent to a carrier about an incident that might lead to a claim; "
      "%d rows, January 2012 to September 2026." %
      counts[K.NOTICES])
    a("")
    a("## Carrier loss runs")
    a("")
    a("%s, %s and %s are the carriers' files as received, in their own layouts and field names. Quillan Mutual insures "
      "Tolliver Falls Hospital and Alderwick Medical Center; Stoneleigh Casualty insures Corvane Medical Group and "
      "Brennick Community Hospital. A run lists every file the carrier holds under the entity's policies, on its "
      "employed clinicians and on the entity itself." % (K.QM_SEP, K.QM_JUN, K.SC_SEP))
    a("")
    a("## %s" % K.SAFETY)
    a("")
    a("Patient safety event reports from the safety reporting system, January 2019 to September 2026, one row per "
      "report; %d rows. Harm level as assigned at the quality team's review." % counts[K.SAFETY])
    a("")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(t))
