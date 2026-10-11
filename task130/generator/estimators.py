"""Every figure the build reports, computed from the shipped files only: the golden tier table and programme
baseline, each rung of the ladder, the renewal letter's exhibits, the rebate statement's counts, and the cost-only
regression the design note compares against."""
import datetime as dt
import os
from collections import defaultdict

import numpy as np
import pandas as pd
from openpyxl import load_workbook

import world as W

TM, TR = W.TREND_MED, W.TREND_RX
MID26 = 15            # months, 1 July 2026 to 1 October 2027
MID_AN = 24           # months, 1 October 2025 to 1 October 2027


def months_between(a, b):
    return (b.year - a.year) * 12 + b.month - a.month


class Files:
    def __init__(self, d):
        self.d = d
        md = pd.read_csv(os.path.join(d, "medical_claim_lines_2025-01_2027-02.csv"), dtype=str, keep_default_na=False)
        for c in ("allowed_amount", "plan_paid_amount", "deductible_amount", "copay_amount", "coinsurance_amount"):
            md[c] = pd.to_numeric(md[c].replace("", "0")).astype(float)
        md["d0"] = pd.to_datetime(md["first_date_of_service"]).dt.date
        md["paid"] = pd.to_datetime(md["paid_date"]).dt.date
        self.med = md
        rx = pd.read_csv(os.path.join(d, "pharmacy_claims_2025-01_2027-02.csv"), dtype=str, keep_default_na=False)
        for c in ("GROSS_AMT_DUE", "PLAN_PAY_AMT", "PATIENT_PAY_AMT"):
            rx[c] = pd.to_numeric(rx[c]).astype(float)
        rx["dos"] = pd.to_datetime(rx["DATE_OF_SERVICE"], format="%Y%m%d").dt.date
        rx["proc"] = pd.to_datetime(rx["DATE_PROCESSED"], format="%Y%m%d").dt.date
        el = pd.read_csv(os.path.join(d, "eligibility_spans_2027-02.csv"), dtype=str, keep_default_na=False)
        self.el = el
        key = {}
        for r in el.itertuples(index=False):
            k = (int(r.employee_number), r.person_code)
            assert key.get(k, r.member_id) == r.member_id, "one employee number and person code, one member"
            key[k] = r.member_id
        rx["member_id"] = [key.get((int(c), p), "") for c, p in zip(rx["CARDHOLDER_ID"], rx["PERSON_CD"])]
        self.rx = rx
        self.spans = defaultdict(list)
        for r in el.itertuples(index=False):
            s = dt.date.fromisoformat(r.span_start)
            e = dt.date.fromisoformat(r.span_end) if r.span_end else dt.date(9999, 12, 31)
            self.spans[r.member_id].append((s, e))
        xw = pd.read_csv(os.path.join(d, "tpa_conversion_crosswalk_2025-01-01.csv"), dtype=str, keep_default_na=False)
        self.xw = {a + "-" + b: m for a, b, m in zip(xw.prior_subscriber_number, xw.prior_member_suffix, xw.member_id)}

    # ---------------------------------------------------------------- the conventions
    def mm(self, member, a, b):
        """member months in [a, b]: months whose first day the member is covered"""
        n, d = 0, dt.date(a.year, a.month, 1)
        if d < a:
            d = W.add_months(d, 1)
        while d <= b:
            if any(s <= d <= e for s, e in self.spans.get(member, ())):
                n += 1
            d = W.add_months(d, 1)
        return n

    def med_net(self, cutoff=None, netted=True, field="allowed_amount"):
        md = self.med if cutoff is None else self.med[self.med.paid <= cutoff]
        if netted:
            gone = set(md.loc[md.claim_frequency_code == "7", "original_claim_number"])
            md = md[~md.claim_number.isin(gone)]
        return md

    def rx_net(self, cutoff=None, reversals=True):
        rx = self.rx if cutoff is None else self.rx[self.rx.proc <= cutoff]
        if not reversals:
            rx = rx[rx.TXN_CD == "B1"]
        return rx

    def allowed(self, members, a, b, cutoff=None, netted=True, paid_field=False, reversals=True):
        """{member: [medical, pharmacy]} for claims incurred in [a, b]"""
        ms = set(members)
        md = self.med_net(cutoff, netted)
        md = md[(md.d0 >= a) & (md.d0 <= b) & md.member_id.isin(ms)]
        rx = self.rx_net(cutoff, reversals)
        rx = rx[(rx.dos >= a) & (rx.dos <= b) & rx.member_id.isin(ms)]
        out = {m: [0.0, 0.0] for m in ms}
        f1 = "plan_paid_amount" if paid_field else "allowed_amount"
        for m, v in md.groupby("member_id")[f1].sum().items():
            out[m][0] = float(v)
        f2 = "PLAN_PAY_AMT" if paid_field else "GROSS_AMT_DUE"
        for m, v in rx.groupby("member_id")[f2].sum().items():
            out[m][1] = float(v)
        return out


def read_list(path):
    wb = load_workbook(path, data_only=True)
    ws = wb["Members"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    return [(r[0], str(r[1]), float(r[5]), int(r[6])) for r in rows if r[0] is not None]


def tier_pmpm(F, groups, a, b, **kw):
    """groups: {tier: [member ids]} -> {tier: (medical PMPM, pharmacy PMPM, member months, medical $, pharmacy $)}"""
    out = {}
    for t, ms in groups.items():
        al = F.allowed(ms, a, b, **kw)
        m = sum(F.mm(x, a, b) for x in ms)
        med = sum(al[x][0] for x in ms)
        rx = sum(al[x][1] for x in ms)
        out[t] = (med / m if m else 0.0, rx / m if m else 0.0, m, med, rx)
    return out


def figures(d):
    """the golden and every rung, from the files in directory d"""
    F = Files(d)
    bid = read_list(os.path.join(d, "northcourt_rfp_2025-07_stratification_2024_claims.xlsx"))
    fin = read_list(os.path.join(d, "northcourt_stratification_2026_claims_final.xlsx"))
    dpath = os.path.join(d, "northcourt_stratification_2026_claims_draft.xlsx")
    dra = read_list(dpath) if os.path.exists(dpath) else []
    bidg, fing, drag = defaultdict(list), defaultdict(list), defaultdict(list)
    for rank, ident, sc, t in bid:
        bidg[t].append(F.xw[ident])
    for rank, ident, sc, t in fin:
        fing[t].append(ident)
    for rank, ident, sc, t in dra:
        drag[t].append(ident)
    T = (1, 2, 3, 4)
    n = {t: len(fing[t]) for t in T}
    N = sum(n.values())
    a26, b26 = dt.date(2026, 1, 1), dt.date(2026, 12, 31)
    an = tier_pmpm(F, bidg, W.ANALOG_START, W.ANALOG_END)
    f26 = tier_pmpm(F, fing, a26, b26)
    G = {}
    G["n"] = n
    G["f26"] = {t: (f26[t][0], f26[t][1]) for t in T}
    G["an"] = {t: (an[t][0], an[t][1]) for t in T}
    G["an_detail"] = {t: (len(bidg[t]), an[t][2], an[t][3], an[t][4]) for t in T}
    G["f26_detail"] = {t: (n[t], f26[t][2], f26[t][3], f26[t][4]) for t in T}
    fm, fr = TM ** (MID_AN / 12), TR ** (MID_AN / 12)
    G["base"] = {t: (an[t][0] * fm, an[t][1] * fr) for t in T}
    G["prog"] = (sum(n[t] * G["base"][t][0] for t in T) / N, sum(n[t] * G["base"][t][1] for t in T) / N)
    G["prog_total"] = G["prog"][0] + G["prog"][1]

    def prog_of(tab):
        return sum(n[t] * (tab[t][0] + tab[t][1]) for t in T) / N

    R = {}
    R["R0"] = {t: (f26[t][0] * TM ** (MID26 / 12), f26[t][1] * TR ** (MID26 / 12)) for t in T}
    c25 = tier_pmpm(F, bidg, dt.date(2025, 1, 1), dt.date(2025, 12, 31))
    R["R2"] = {t: (c25[t][0] * TM ** (27 / 12), c25[t][1] * TR ** (27 / 12)) for t in T}
    R["R3a"] = {t: (an[t][0], an[t][1]) for t in T}
    R["R3b"] = {t: (an[t][0] * TM ** 1.5, an[t][1] * TR ** 1.5) for t in T}
    R["R4a"] = {t: (an[t][3] / (12 * len(bidg[t])) * fm, an[t][4] / (12 * len(bidg[t])) * fr) for t in T}
    apr = {t: [m for m in bidg[t] if F.mm(m, dt.date(2025, 4, 1), dt.date(2025, 4, 30)) == 1] for t in T}
    R["R4b"] = {t: (an[t][3] / (12 * len(apr[t])) * fm, an[t][4] / (12 * len(apr[t])) * fr) for t in T}
    G["r4b_n"] = {t: len(apr[t]) for t in T}
    paid = tier_pmpm(F, bidg, W.ANALOG_START, W.ANALOG_END, paid_field=True)
    R["R6a"] = {t: (paid[t][0] * fm, paid[t][1] * fr) for t in T}
    raw = tier_pmpm(F, bidg, W.ANALOG_START, W.ANALOG_END, netted=False, reversals=False)
    R["R6b"] = {t: (raw[t][0] * fm, raw[t][1] * fr) for t in T}
    rawp = tier_pmpm(F, bidg, W.ANALOG_START, W.ANALOG_END, netted=False, reversals=False, paid_field=True)
    R["R6c"] = {t: (rawp[t][0] * fm, rawp[t][1] * fr) for t in T}
    nd = {t: len(drag[t]) for t in T}
    rk = {}
    rk["R0"] = prog_of(R["R0"])
    for k in ("R2", "R3a", "R3b", "R4a", "R4b", "R6a", "R6b", "R6c"):
        rk[k] = prog_of(R[k])
    rk["R5b"] = sum(nd[t] * (G["base"][t][0] + G["base"][t][1]) for t in T) / sum(nd.values()) if dra else float("nan")
    rk["R5a"] = rk["R0"]
    # pooled: the analog over all 140 bid members, ignoring tiers
    allb = [m for t in T for m in bidg[t]]
    pool = tier_pmpm(F, {0: allb}, W.ANALOG_START, W.ANALOG_END)[0]
    rk["pooled"] = pool[0] * fm + pool[1] * fr
    LF = letter_figures(F)
    rk["R1"] = LF["proj_med"] + LF["proj_rx"]
    # R0 with the 2026 claims completed at the renewal letter's factors
    cf = LF["cf"]
    comp = {}
    for t in T:
        med_c = 0.0
        for mth in range(1, 13):
            a, b = dt.date(2026, mth, 1), W.month_end(W.month_index(dt.date(2026, mth, 1)))
            al = F.allowed(fing[t], a, b)
            med_c += sum(v[0] for v in al.values()) / cf[mth]
        comp[t] = (med_c / f26[t][2] * TM ** (MID26 / 12), f26[t][1] * TR ** (MID26 / 12))
    R["R0c"] = comp
    rk["R0c"] = prog_of(comp)
    rk["regression"], R["regression"] = regression_alternative(F, fing, n)
    G["rungs"] = rk
    G["rung_tiers"] = R
    G["letter"] = LF
    G["bid_tiers"] = {t: len(bidg[t]) for t in T}
    G["draft_n"] = nd
    G["members"] = {"bid": bidg, "final": fing, "draft": drag}
    G["zero26"] = sum(1 for t in T for m in fing[t] if F.mm(m, a26, b26) == 0)
    return G, F


def regression_alternative(F, fing, n):
    """a cost-only regression of 2026 allowed on 2025 allowed across members enrolled all of 2025 and 2026, applied to
    the enrolled members' 2026 claims and projected three months at trend (calendar 2027 to the programme year)"""
    a25, b25, a26, b26 = dt.date(2025, 1, 1), dt.date(2025, 12, 31), dt.date(2026, 1, 1), dt.date(2026, 12, 31)
    full = [m for m in F.spans if F.mm(m, a25, b25) == 12 and F.mm(m, a26, b26) == 12]
    x = F.allowed(full, a25, b25)
    y = F.allowed(full, a26, b26)
    coef = []
    for k in (0, 1):
        X = np.array([x[m][k] for m in full])
        Y = np.array([y[m][k] for m in full])
        b = float(np.cov(X, Y, bias=True)[0, 1] / np.var(X))
        a = float(Y.mean() - b * X.mean())
        coef.append((a, b))
    out = {}
    for t, ms in fing.items():
        al = F.allowed(ms, a26, b26)
        pm = [0.0, 0.0]
        for m in ms:
            mm26 = F.mm(m, a26, b26)
            for k in (0, 1):
                ann = al[m][k] * 12 / mm26 if mm26 else 0.0
                pm[k] += (coef[k][0] + coef[k][1] * ann) / 12
        out[t] = (pm[0] / len(ms) * TM ** 0.25, pm[1] / len(ms) * TR ** 0.25)
    N = sum(n.values())
    return sum(n[t] * (out[t][0] + out[t][1]) for t in out) / N, out


# ------------------------------------------------------------------ the renewal letter's exhibits
def letter_figures(F):
    cut = W.LETTER_PAID
    allm = list(F.spans)
    mm = {y: sum(F.mm(m, dt.date(y, 1, 1), dt.date(y, 12, 31)) for m in allm) for y in (2025, 2026)}
    md = F.med_net(cut)
    rx = F.rx_net(cut)
    med_paid = {y: float(md[(md.d0 >= dt.date(y, 1, 1)) & (md.d0 <= dt.date(y, 12, 31))].allowed_amount.sum()) for y in (2025, 2026)}
    rxa = {y: float(rx[(rx.dos >= dt.date(y, 1, 1)) & (rx.dos <= dt.date(y, 12, 31))].GROSS_AMT_DUE.sum()) for y in (2025, 2026)}
    # completion factors from the second half of 2025: share of the final allowed paid k months after the month
    final = F.med_net(None)
    cfk = {}
    for k in range(1, 7):
        num = den = 0.0
        for mth in range(7, 13):
            a = dt.date(2025, mth, 1)
            b = W.month_end(W.month_index(a))
            cut_k = W.month_end(W.month_index(a) + k)
            tot = final[(final.d0 >= a) & (final.d0 <= b)].allowed_amount.sum()
            part = F.med_net(cut_k)
            part = part[(part.d0 >= a) & (part.d0 <= b)].allowed_amount.sum()
            num += part
            den += tot
        cfk[k] = round(num / den, 3)
    cf = {m: 1.0 for m in range(1, 13)}
    for mth in range(7, 13):
        cf[mth] = cfk[13 - mth]
    by_month = {}
    for mth in range(1, 13):
        a = dt.date(2026, mth, 1)
        b = W.month_end(W.month_index(a))
        by_month[mth] = float(md[(md.d0 >= a) & (md.d0 <= b)].allowed_amount.sum())
    med_comp = {2025: med_paid[2025], 2026: sum(by_month[m] / cf[m] for m in range(1, 13))}
    proj_med = med_comp[2026] / mm[2026] * TM ** 1.5
    proj_rx = rxa[2026] / mm[2026] * TR ** 1.5
    # funding: the current rates were set last year; the increase is the projected cost over the current revenue
    el = F.el
    feb = dt.date(2027, 2, 1)
    enrol = defaultdict(int)
    members = 0
    for r in el.itertuples(index=False):
        s = dt.date.fromisoformat(r.span_start)
        e = dt.date.fromisoformat(r.span_end) if r.span_end else dt.date(9999, 12, 31)
        if s <= feb <= e:
            members += 1
            if r.relationship == "SUB":
                enrol[("PPO" if r.plan_code == "RCV-PPO" else "HDHP", r.coverage_tier)] += 1
    subs = sum(enrol.values())
    now = RATES_NOW
    revenue = sum(enrol[(p, t)] * now[p][t] for p, t in enrol)
    cost = (proj_med + proj_rx) * members * 1.03 + subs * (38.75 + 78.20 + 4.15)
    incr = cost / revenue - 1
    new = {p: {t: round(now[p][t] * (1 + incr), 2) for t in now[p]} for p in now}
    return {"mm": mm, "med_paid": med_paid, "med_comp": med_comp, "rx": rxa, "cf": cf, "cfk": cfk,
            "proj_med": proj_med, "proj_rx": proj_rx, "rate_incr": round(100 * incr, 1), "rates_now": now, "rates_new": new,
            "subs_feb": subs, "members_feb": members, "enrol": dict(enrol)}


RATES_NOW = {"PPO": {"EE": 1058.40, "ES": 2201.30, "EC": 1873.60, "FAM": 3002.70},
             "HDHP": {"EE": 861.20, "ES": 1791.10, "EC": 1524.40, "FAM": 2443.10}}


# ------------------------------------------------------------------ the rebate statement
G_BRAND, G_SPEC = 212.50, 1640.00


def rebate_figures(F):
    rx = F.rx[(F.rx.dos >= dt.date(2026, 1, 1)) & (F.rx.dos <= dt.date(2026, 12, 31)) & (F.rx.BRAND_GENERIC_CD == "B")]
    rx = rx.assign(q=[(d.month - 1) // 3 + 1 for d in rx.dos], sg=np.where(rx.TXN_CD == "B2", -1, 1),
                   spec=rx.PHARMACY_NCPDP_ID == W.NCPDP_IDS[W.SPECIALTY_PHARMACY])
    brand = {q: int(rx[(rx.q == q) & ~rx.spec].sg.sum()) for q in (1, 2, 3, 4)}
    spec = {q: int(rx[(rx.q == q) & rx.spec].sg.sum()) for q in (1, 2, 3, 4)}
    by = []
    for (ndc, name, sp), g in rx.groupby(["NDC", "DRUG_NAME", "spec"]):
        qs = [int(g[g.q == q].sg.sum()) for q in (1, 2, 3, 4)]
        by.append((name, ndc, bool(sp), qs))
    by.sort(key=lambda x: (not x[2], x[0]))
    return {"brand": brand, "spec": spec, "g_brand": G_BRAND, "g_spec": G_SPEC, "by_drug": by}


def stoploss_census(F):
    jan = dt.date(2027, 1, 1)
    n = 0
    for r in F.el.itertuples(index=False):
        s = dt.date.fromisoformat(r.span_start)
        e = dt.date.fromisoformat(r.span_end) if r.span_end else dt.date(9999, 12, 31)
        if r.relationship == "SUB" and s <= jan <= e:
            n += 1
    return n
