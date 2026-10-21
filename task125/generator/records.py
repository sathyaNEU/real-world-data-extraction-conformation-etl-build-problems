"""Write the HLS data files: the person-quarter microdata, the wave-1 call records (the intakes the call-record app
covers), the address register, the population controls, the release tables to the latest published quarter, and the
two files the analysis never needs (the benefit claimant counts and the field staffing list)."""
import csv
import datetime as dt
import json

import numpy as np
import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl.styles import Alignment, Font

import models as M
import world as W

PERSON = "hls_person_quarter_2023q1_2026q3.parquet"
CALLS = "hls_call_records_2025q1_2026q3.csv"
REGISTER = "hls_address_register_2025q1_2026q3.csv"
CONTROLS = "population_controls_2023q1_2027q4.xlsx"
RELEASE = "hls_release_tables_2026q2.xlsx"
CLAIMANTS = "jobseeker_support_claimants_2023-01_2026-09.csv"
STAFFING = "field_staffing_2026-10.json"

CHECK_FROM = (2026, 3)            # serials carry a check digit from the July to September 2026 processing run
REPLAY = (dt.datetime(2026, 8, 4), dt.datetime(2026, 9, 25, 23, 59))
REPLAY_SHARE = 0.11
SAME_DAY = 0.11


def _rng(*k):
    return np.random.default_rng([W.SEED, 7] + list(k))


def serial(c, i):
    return "H%02d%d%05d" % (c[0] % 100, c[1], i + 1)


def check_digit(s):
    digits = [int(ch) for ch in s[1:]]
    return str(sum((k + 2) * d for k, d in enumerate(reversed(digits))) % 11 % 10)


def as_filed(s, q):
    return s + check_digit(s) if W.qindex(q) >= W.qindex(CHECK_FROM) else s


def area_code(region, district):
    return "%d%02d" % (W.RCODES.index(W.RCODES[region]) + 1, district)


def hmm(h):
    """Telephone interviewing stores hours and minutes as h.mm."""
    whole = int(h)
    return round(whole + round((h - whole) * 60) / 100, 2)


def quarter_rows(intakes, q):
    rows = []
    for w in range(1, W.WAVES + 1):
        c = W.qnext(q, -(w - 1))
        d = intakes.get(c)
        if d is None:
            continue
        hh, pres, st = d["hh"], d["present"][w - 1], d["statuses"][w - 1]
        for h in np.nonzero(d["resp"][w - 1])[0]:
            members = np.nonzero((hh == h) & pres)[0]
            if w > 1:                                   # the telephone system lists the household oldest first
                members = sorted(members, key=lambda i: (d["dob"][i], i))
            for pno, i in enumerate(members, start=1):
                hrs = d["hours"][w - 1][i] if st[i] == 0 else None
                if hrs is not None and w > 1:
                    hrs = hmm(hrs)
                rows.append(dict(quarter=W.qlabel(q), serial=as_filed(serial(c, h), q), person_no=pno, wave=w,
                                 mode="CAPI" if w == 1 else "CATI", region=W.RCODES[d["region"][h]],
                                 sex=int(d["sex"][i]), date_of_birth=d["dob"][i], age=int(d["ages_q"][w - 1][i]),
                                 ilo_status=int(st[i]) + 1, actual_hours=hrs, _hh=(c, int(h)), _i=int(i)))
    s = M.quarter_sample(intakes, q)
    assert len(s["status"]) == len(rows)
    w = M.rake(s["region"], s["cell"], q)
    weight = {(int(c), int(i)): float(x) for c, i, x in zip(s["intake"], s["person"], w)}
    for r in rows:
        r["weight"] = round(weight[(W.qindex(r["_hh"][0]), r["_i"])], 2)
    return rows


def person_file(intakes):
    rows = []
    for q in W.qrange(W.FIRST_Q, W.LAST_Q):
        rows += quarter_rows(intakes, q)
    rows.sort(key=lambda r: (r["quarter"], r["serial"], r["person_no"]))
    return rows


def write_person(path, rows):
    cols = ["quarter", "serial", "person_no", "wave", "mode", "region", "sex", "date_of_birth", "age", "ilo_status",
            "actual_hours", "weight"]
    schema = pa.schema([("quarter", pa.string()), ("serial", pa.string()), ("person_no", pa.int16()), ("wave", pa.int8()),
                        ("mode", pa.string()), ("region", pa.string()), ("sex", pa.int8()), ("date_of_birth", pa.date32()),
                        ("age", pa.int16()), ("ilo_status", pa.int8()), ("actual_hours", pa.float64()),
                        ("weight", pa.float64())])
    t = pa.Table.from_pylist([{k: r[k] for k in cols} for r in rows], schema=schema)
    pq.write_table(t, path, compression="zstd", row_group_size=100_000, write_statistics=True)
    return len(rows)


def calls(intakes):
    """Every visit to every address issued in the intakes the call-record app covers, as it uploaded them."""
    rows = []
    for c in W.qrange(W.CALLS_FROM, W.LAST_Q):
        d = intakes[c]
        rng = _rng(*c)
        start = dt.datetime(c[0], 3 * c[1] - 2, 1)
        n = len(d["region"])
        for h in range(n):
            area = area_code(d["region"][h], d["district"][h])
            interviewer = "I%s%d" % (area, 1 + (h % 3))
            t = start + dt.timedelta(days=int(rng.integers(1, 48)), hours=int(rng.integers(9, 20)),
                                     minutes=int(rng.choice([0, 10, 20, 30, 40, 50])))
            if not d["eligible"][h]:
                seq = [(1, t, 410)]
            else:
                need, maxv = int(d["need"][h]), d["maxv"]
                seq = []
                for v in range(1, min(need, maxv) + 1):
                    out = 310 if v < need else (210 if d["refuse"][h] else 110)
                    seq.append((v, t, out))
                    if rng.random() < SAME_DAY and t.hour < 17:
                        t = t.replace(hour=int(rng.integers(17, 20)), minute=int(rng.choice([0, 15, 30, 45])))
                    else:
                        t = t + dt.timedelta(days=int(rng.integers(2, 10)))
                        t = t.replace(hour=int(rng.integers(9, 20)), minute=int(rng.choice([0, 10, 20, 30, 40, 50])))
            for v, tv, out in seq:
                up = tv + dt.timedelta(minutes=int(rng.integers(40, 30 * 60)))
                base = dict(serial=as_filed(serial(c, h), c), intake=W.qlabel(c), area_code=area, interviewer=interviewer,
                            visit_no=v, visit_start=tv, outcome=out, uploaded=up)
                rows.append(base)
                if c == W.RELEASE_Q and out != 110 and REPLAY[0] <= tv <= REPLAY[1] and rng.random() < REPLAY_SHARE:
                    rows.append(dict(base, uploaded=up + dt.timedelta(hours=int(rng.integers(20, 70))), _replay=True))
    rows.sort(key=lambda r: (r["uploaded"], r["serial"], r["visit_no"]))
    for k, r in enumerate(rows):
        r["record_id"] = "CR%07d" % (1_204_311 + k)
    return rows


def write_calls(path, rows):
    cols = ["record_id", "serial", "intake", "area_code", "interviewer", "visit_no", "visit_start", "outcome", "uploaded"]
    with open(path, "w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh)
        wr.writerow(cols)
        for r in sorted(rows, key=lambda r: r["record_id"]):
            wr.writerow([r[k].strftime("%Y-%m-%d %H:%M") if isinstance(r[k], dt.datetime) else r[k] for k in cols])
    return len(rows)


def write_register(path, intakes):
    n = 0
    with open(path, "w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh)
        wr.writerow(["serial", "intake", "area_code", "region", "urban_rural", "dwelling_type"])
        for c in W.qrange(W.CALLS_FROM, W.LAST_Q):
            d = intakes[c]
            rng = _rng(99, *c)
            for h in range(len(d["region"])):
                r = d["region"][h]
                urban = "urban" if rng.random() < [0.92, 0.71, 0.58, 0.49, 0.44, 0.36][r] else "rural"
                dwell = rng.choice(["detached house", "semi-detached or terraced", "flat or apartment", "other"],
                                   p=[0.55, 0.18, 0.24, 0.03])
                wr.writerow([serial(c, h), W.qlabel(c), area_code(r, d["district"][h]), W.RCODES[r], urban, dwell])
                n += 1
    return n


def write_controls(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Region"
    ws.append(["quarter", "basis"] + ["%s %s" % (c, W.RNAME[c]) for c in W.RCODES] + ["total"])
    ws2 = wb.create_sheet("Age and sex")
    hdr = ["quarter", "basis"]
    for g in W.AGE_GROUPS:
        hdr += ["%s male" % g, "%s female" % g]
    ws2.append(hdr + ["total"])
    for q in W.qrange(W.FIRST_Q, (2027, 4)):
        reg, cells = M.controls(q)
        basis = "2025-based estimates and projections" if W.qindex(q) >= W.qindex((2026, 1)) else "2023-based projections"
        ws.append([W.qlabel(q), basis] + [int(x) for x in reg] + [int(reg.sum())])
        ws2.append([W.qlabel(q), basis] + [int(x) for x in cells] + [int(cells.sum())])
    for s in (ws, ws2):
        for cell in s[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        s.freeze_panes = "C2"
        s.column_dimensions["A"].width = 9
        s.column_dimensions["B"].width = 34
        for col in range(3, s.max_column + 1):
            s.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 13
            for row in range(2, s.max_row + 1):
                s.cell(row, col).number_format = "#,##0"
    wn = wb.create_sheet("Notes")
    for line in ["Population aged 15 and over usually resident in Merrowland, persons, at the middle of each quarter.",
                 "From January to March 2026 the controls are on the 2025-based estimates; earlier quarters stay on the 2023-based projections.",
                 "Quarters after July to September 2026 are projections.",
                 "Source: Merrowland Statistics, Population Estimates and Projections."]:
        wn.append([line])
    wn.column_dimensions["A"].width = 110
    wb.save(path)


def decimal_hours(h, mode):
    if mode != "CATI":
        return h
    whole = int(h)
    return whole + round((h - whole) * 100) / 60


def write_release(path, intakes, person_rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Table 1"
    ws.append(["Household Labour Survey: headline rates, people aged 15 and over (not seasonally adjusted)"])
    ws.append([])
    ws.append(["Quarter", "Unemployment rate (%)", "Employment rate (%)", "Unemployed (000)", "Employed (000)",
               "Total actual weekly hours (000)"])
    hours = {}
    for p in person_rows:
        if p["ilo_status"] == 1:
            hours[p["quarter"]] = hours.get(p["quarter"], 0.0) + p["weight"] * decimal_hours(p["actual_hours"], p["mode"])
    regional = []
    for q in W.qrange(W.FIRST_Q, W.LAST_PUBLISHED):
        s = M.quarter_sample(intakes, q)
        w = M.rake(s["region"], s["cell"], q)
        u, e = M.rates(s["status"], w)
        ws.append([W.qname(q), round(u, 1), round(e, 1), round(w[s["status"] == 1].sum() / 1000, 1),
                   round(w[s["status"] == 0].sum() / 1000, 1), round(hours[W.qlabel(q)] / 1000, 1)])
        regional.append([W.qname(q)] + [round(M.rates(s["status"][s["region"] == r], w[s["region"] == r])[0], 1)
                                        for r in range(len(W.RCODES))])
    ws2 = wb.create_sheet("Table 4")
    ws2.append(["Household Labour Survey: unemployment rate by region (%), people aged 15 and over"])
    ws2.append([])
    ws2.append(["Quarter"] + [W.RNAME[c] for c in W.RCODES])
    for r in regional:
        ws2.append(r)
    for s, widths in ((ws, [24, 14, 14, 14, 14, 16]), (ws2, [24] + [12] * 6)):
        s["A1"].font = Font(bold=True)
        for cell in s[3]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for k, wdt in enumerate(widths, start=1):
            s.column_dimensions[openpyxl.utils.get_column_letter(k)].width = wdt
        for row in s.iter_rows(min_row=4):
            for cell in row[1:]:
                cell.number_format = "0.0"
        s.freeze_panes = "B4"
    wn = wb.create_sheet("Notes")
    for line in ["Estimates for April to June 2026, released 13 August 2026.",
                 "Rates are of people aged 15 and over: unemployed as a share of the labour force, employed as a share of the population.",
                 "Estimates are weighted to the population controls in force for each quarter; see the HLS methods note.",
                 "Total actual weekly hours: hours actually worked in the reference week by employed people, all jobs.",
                 "Figures are rounded independently."]:
        wn.append([line])
    wn.column_dimensions["A"].width = 110
    wb.save(path)


def write_claimants(path):
    rng = _rng(31)
    base = {"VEL": 41_800, "CAR": 22_600, "PEL": 18_900, "HAS": 15_200, "TIR": 14_100, "DUN": 11_300}
    n = 0
    with open(path, "w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh)
        wr.writerow(["month", "region", "jobseeker_support_work_ready", "jobseeker_support_health_condition"])
        for y in range(2023, 2027):
            for m in range(1, 13):
                if (y, m) > (2026, 9):
                    break
                t = (y - 2023) * 12 + m
                for code in W.RCODES:
                    level = base[code] * (1 + 0.004 * t + 0.02 * np.sin(2 * np.pi * (m - 2) / 12))
                    if (y, m) >= (2025, 4):         # the 2025 work-capacity reassessment moved claimants between groups
                        level *= 1.09
                    ready = int(level * (1 + rng.normal(0, 0.006)))
                    health = int(level * 0.62 * (1 + rng.normal(0, 0.006)) * (0.90 if (y, m) >= (2025, 4) else 1))
                    wr.writerow(["%d-%02d" % (y, m), code, ready, health])
                    n += 1
    return n


def write_staffing(path):
    rng = _rng(41)
    out = {"office": W.OFFICE, "as_at": "2026-10-01", "unit": "Field Operations",
           "note": "Field interviewers on the face-to-face panel by region. FTE as contracted for October to December 2026.",
           "regions": []}
    for code in W.RCODES:
        k = W.DISTRICTS[code] * 3
        staff = []
        for i in range(k):
            staff.append({"interviewer": "I%d%02d%d" % (W.RCODES.index(code) + 1, i // 3 + 1, 1 + i % 3),
                          "fte": float(rng.choice([0.6, 0.8, 1.0], p=[0.3, 0.3, 0.4])),
                          "contract": str(rng.choice(["permanent", "fixed term"], p=[0.8, 0.2]))})
        out["regions"].append({"region": code, "name": W.RNAME[code], "interviewers": staff})
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(out, fh, indent=2)
        fh.write("\n")
