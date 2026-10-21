"""The goldens, computed from target/ only: the methods note to the Government Statistician
(hls_visit_limit_note_2026q3.docx, with its chart) and the gap workbook (hls_visit_limit_path_2027.xlsx), and
generator/figures.json with every figure the write-up states."""
import io
import json
import os
import sys

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt   # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import analysis as AN      # noqa: E402
import asks as A           # noqa: E402
import postprocess as PP   # noqa: E402
import records as R        # noqa: E402
import world as W          # noqa: E402

DOCX = "hls_visit_limit_note_2026q3.docx"
XLSX = "hls_visit_limit_path_2027.xlsx"
REF = "MS-LMS-2026-031"
TO = ("Shane Turner", "Government Statistician")
INK, INK2, GRID, PAPER = "#1d2329", "#56606a", "#dfe3e7", "#ffffff"
SERIES_C, CAP_C, GAP_C, GAP_C2 = "#2f5d8a", "#b5401f", "#b5401f", "#e3a48f"
SHORT = {1: "Jan-Mar", 2: "Apr-Jun", 3: "Jul-Sep", 4: "Oct-Dec"}


def short(q):
    return "%s %d" % (SHORT[q[1]], q[0])


def r2(x):
    return float("%.2f" % (round(x * 100 + 1e-9) / 100)) if x >= 0 else -float("%.2f" % (round(-x * 100 + 1e-9) / 100))


def f2(x):
    return "%.2f" % r2(x)


def compute(target):
    d = AN.load(target)
    h = AN.headline(d)
    rows = []
    for r in d["person"].itertuples(index=False):
        rows.append(dict(quarter=r.quarter, serial=r.serial, person_no=r.person_no, mode=r.mode, region=r.region,
                         sex=r.sex, date_of_birth=r.date_of_birth, ilo_status=r.ilo_status, actual_hours=r.actual_hours,
                         weight=r.weight))
    import csv
    import datetime as dt
    with open(os.path.join(target, R.CALLS), encoding="utf-8") as fh:
        calls = [dict(r, visit_no=int(r["visit_no"]), visit_start=dt.datetime.strptime(r["visit_start"], "%Y-%m-%d %H:%M"))
                 for r in csv.DictReader(fh)]
    with open(os.path.join(target, R.REGISTER), encoding="utf-8") as fh:
        register = list(csv.DictReader(fh))
    asks = dict(visits=A.visits(calls, register), flows=A.flows(rows), hours=A.hours(rows))
    return dict(d=d, h=h, asks=asks)


def figures(c):
    h, a = c["h"], c["asks"]
    path = h["path"]
    return {
        "rise_unemployment_rate_points": f2(h["rise_u"]),
        "cap_share_of_rise_points": f2(h["cap_u"]),
        "labour_market_share_of_rise_points": f2(h["real_u"]),
        "published_unemployment_rate_2026Q2": f2(h["pub_q2"][0]),
        "published_unemployment_rate_2026Q3": f2(h["pub_q3"][0]),
        "unemployment_rate_2026Q3_without_cap": f2(h["without_u"]),
        "unemployment_rate_2026Q2_without_cap": f2(h["without_q2_u"]),
        "unemployment_rate_gap_2026Q2_points": f2(h["gap_prior"][0]),
        "change_employment_rate_points": f2(h["rise_e"]),
        "cap_share_of_employment_change_points": f2(h["cap_e"]),
        "cap_took_off_employment_change_points": f2(-h["cap_e"]),
        "labour_market_share_of_employment_change_points": f2(h["real_e"]),
        "published_employment_rate_2026Q2": f2(h["pub_q2"][1]),
        "published_employment_rate_2026Q3": f2(h["pub_q3"][1]),
        "employment_rate_2026Q3_without_cap": f2(h["without_e"]),
        "employment_rate_2026Q2_without_cap": f2(h["without_q2_e"]),
        "employment_rate_gap_2026Q2_points": f2(h["gap_prior"][1]),
        "employment_rate_below_six_visit_rate_2026Q2_points": f2(-h["gap_prior"][1]),
        "unemployment_rate_gap_points": {short(q): f2(g[0]) for q, g in zip(AN.FORWARD, path)},
        "employment_rate_gap_points": {short(q): f2(g[1]) for q, g in zip(AN.FORWARD, path)},
        "employment_rate_below_six_visit_rate_points": {short(q): f2(-g[1]) for q, g in zip(AN.FORWARD, path)},
        "employment_fall_points": f2(-h["rise_e"]),
        "employment_rise_without_cap_points": f2(h["real_e"]),
        "template_quarter": short(W.TEMPLATE_Q),
        "template_quarter_name": W.qname(W.TEMPLATE_Q),
        "clean_from_quarter": short(W.qnext(W.LIMIT_TO, W.WAVES)),
        "clean_from_quarter_name": W.qname(W.qnext(W.LIMIT_TO, W.WAVES)),
        "six_visits_from_quarter": short(W.qnext(W.LIMIT_TO)),
        "gap_tolerance_points": {"unemployment": "0.05", "employment": "0.15"},
        "first_interview_visits_2026Q3_intake": {W.RNAME[g]: a["visits"]["golden"][g] for g in W.RCODES},
        "unemployed_to_employed_share_pct": {W.RNAME[g]: f2(a["flows"]["golden"][g]) for g in W.RCODES},
        "unemployed_to_employed_counts": {W.RNAME[g]: list(a["flows"]["golden_n"][g]) for g in W.RCODES},
        "total_weekly_hours_thousand": {W.RNAME[g]: "%.1f" % a["hours"]["golden"][g] for g in W.RCODES},
    }


def chart_png(c):
    h = c["h"]
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 8.6, "svg.hashsalt": "t125",
                         "axes.edgecolor": INK2, "axes.labelcolor": INK2, "xtick.color": INK2, "ytick.color": INK2})
    fig, (ax, bx) = plt.subplots(1, 2, figsize=(7.4, 3.35), dpi=170, gridspec_kw={"width_ratios": [1.55, 1]})
    fig.patch.set_facecolor(PAPER)
    qs = [q for q, _ in h["series"]]
    ys = [r[0] for _, r in h["series"]]
    xs = list(range(len(qs)))
    ax.plot(xs, ys, color=SERIES_C, linewidth=1.8, marker="o", markersize=3.2, zorder=3, label="Published")
    wo = [(x, v) for x, q, v in zip(xs, qs, ys) if q in (W.PRIOR_Q, W.RELEASE_Q)]
    wox = [x for x, _ in wo]
    woy = [h["without_q2_u"], h["without_u"]]
    ax.plot([xs[xs.index(wox[0]) - 1]] + wox, [ys[wox[0] - 1]] + woy, color=CAP_C, linewidth=1.4,
            linestyle=(0, (3, 2)), zorder=3)
    ax.scatter(wox, woy, s=26, color=PAPER, edgecolor=CAP_C, linewidth=1.4, zorder=4, label="Six visits throughout")
    ax.annotate("%.2f" % r2(ys[-1]), (xs[-1], ys[-1]), xytext=(5, 2), textcoords="offset points", fontsize=7.6,
                color=SERIES_C)
    ax.annotate("%.2f" % r2(h["without_u"]), (xs[-1], h["without_u"]), xytext=(5, -9), textcoords="offset points",
                fontsize=7.6, color=CAP_C)
    ax.set_xticks([x for x, q in zip(xs, qs) if q[1] == 1])
    ax.set_xticklabels(["Jan-Mar\n%d" % q[0] for q in qs if q[1] == 1])
    lo = min(min(ys), min(woy)) - 0.22
    hi = max(max(ys), max(woy)) + 0.22
    ax.set_ylim(round(lo * 5) / 5, round(hi * 5) / 5)
    ax.set_ylabel("Unemployment rate, %")
    ax.yaxis.grid(True, color=GRID, linewidth=0.7, zorder=0)
    ax.legend(frameon=False, fontsize=7.6, loc="upper left")
    ax.set_title("Published rate, and the rate six visits throughout would have given", fontsize=8.6, color=INK,
                 loc="left")
    gaps = [g[0] for g in h["path"]]
    labels = [short(q).replace(" ", "\n") for q in AN.FORWARD]
    cols = [GAP_C if W.capped(q) else GAP_C2 for q in AN.FORWARD]
    bx.bar(range(5), gaps, color=cols, width=0.62, zorder=3)
    for i, g in enumerate(gaps):
        bx.text(i, g + 0.008, "%.2f" % r2(g), ha="center", va="bottom", fontsize=7.6, color=INK)
    cut = sum(1 for q in AN.FORWARD if W.capped(q)) - 0.5
    top = max(gaps) * 1.30
    bx.axvline(cut, color=INK2, linewidth=0.8, linestyle=":")
    bx.text(cut + 0.06, top * 0.98, "six visits from\n%s" % short(W.qnext(W.LIMIT_TO)), fontsize=7, color=INK2,
            va="top")
    bx.set_xticks(range(5))
    bx.set_xticklabels(labels, fontsize=7.4)
    bx.set_ylim(0, top)
    bx.set_ylabel("Points above the six-visit rate")
    bx.yaxis.grid(True, color=GRID, linewidth=0.7, zorder=0)
    bx.set_title("Projected gap in the published rate", fontsize=8.6, color=INK, loc="left")
    for a_ in (ax, bx):
        a_.set_facecolor(PAPER)
        for s in ("top", "right"):
            a_.spines[s].set_visible(False)
        a_.tick_params(length=0)
    fig.suptitle("The cap builds as the capped intakes spread through the panel, and is most of the July to "
                 "September rise", x=0.012, ha="left", fontsize=9.8, fontweight="bold", color=INK)
    fig.tight_layout(rect=(0, 0, 1, 0.965), w_pad=2.2)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor=PAPER, metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def write_docx(path, c, fig):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    h = c["h"]
    path_ = h["path"]
    head_name, head_title = W.PEOPLE["head"]
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.9)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    st.paragraph_format.space_after = Pt(6)

    def widths(table, cms):
        table.autofit = False
        grid = table._tbl.tblGrid
        for gc, w_ in zip(grid.findall(qn("w:gridCol")), cms):
            gc.set(qn("w:w"), str(int(w_ * 567)))
        for row in table.rows:
            for cell, w_ in zip(row.cells, cms):
                cell.width = Cm(w_)

    def para(text="", bold=False, size=None, color=None, after=None, italic=False, keep=False):
        p = doc.add_paragraph()
        if text:
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        if keep:
            p.paragraph_format.keep_with_next = True
        return p

    para("Merrowland Statistics  |  Labour Market Statistics", size=8.5, color="56606A", after=8)
    head = doc.add_table(rows=5, cols=2)
    for i, (k, v) in enumerate([("To", "%s, %s" % TO), ("From", "%s, %s" % (head_name, head_title)),
                                ("Date", "23 October 2026"), ("Reference", REF),
                                ("Subject", "HLS July to September 2026: what the three-visit cap does to the published rates")]):
        r = head.cell(i, 0).paragraphs[0].add_run(k)
        r.bold = True
        r.font.size = Pt(9)
        r2_ = head.cell(i, 1).paragraphs[0].add_run(v)
        r2_.font.size = Pt(9)
        for cell in head.rows[i].cells:
            cell.paragraphs[0].paragraph_format.space_after = Pt(1)
    widths(head, [2.6, 13.8])
    para(after=2)
    g = [f2(x[0]) for x in path_]
    e = [f2(x[1]) for x in path_]
    de = h["real_e"]
    gp = h["gap_prior"]
    para("Most of the July to September rise in unemployment is the first-interview visit cap, whose effect on the "
         "published rates goes on building into 2027 before it fades", bold=True, size=12.5, after=8)
    para("The published unemployment rate rose %s point to %s per cent between April to June and July to September. "
         "The three-visit cap on first interviews accounts for %s point of that rise and the labour market for %s "
         "point. The cap was already in the April to June rate: on six visits throughout, the rate would have been "
         "%s per cent in April to June and %s per cent in July to September. The employment rate fell %s point to %s "
         "per cent; the cap took %s point off the change, so on six visits throughout employment rose %s point, from "
         "%s to %s per cent."
         % (f2(h["rise_u"]), f2(h["pub_q3"][0]), f2(h["cap_u"]), f2(h["real_u"]), f2(h["without_q2_u"]),
            f2(h["without_u"]), f2(-h["rise_e"]), f2(h["pub_q3"][1]), f2(-h["cap_e"]), f2(de),
            f2(h["without_q2_e"]), f2(h["without_e"])))
    para("The cap removes the households we find only at a fourth, fifth or sixth visit: adults out by day, most of "
         "them in work. We recruit a household once and keep it for five quarters, so an intake issued under the cap "
         "is short of those households in every quarter it stays in the sample, not only in the quarter it is first "
         "interviewed. First interviews have been capped since the April to June intake, so two of July to "
         "September's five intakes were recruited under the cap, three of October to December's and four of January "
         "to March 2027's. That is what makes the gap grow: it is %s point on the unemployment rate in July to "
         "September 2026, %s in October to December and %s in January to March 2027, and %s point on the employment "
         "rate in the same quarters (%s in April to June 2026)."
         % (g[0], g[1], g[2], " and ".join(", ".join(f2(-x[1]) for x in path_[:3]).rsplit(", ", 1)), f2(-gp[1])))
    para("With first interviews back on six visits from 1 April 2027 the gap falls as the capped intakes leave the "
         "panel, but it does not go while they are in it: the unemployment rate sits %s point above the six-visit "
         "rate in April to June 2027 and %s in July to September 2027, and the employment rate %s and %s point below "
         "it. The published rates are back on a six-visit basis from %s, once the January to March 2027 intake has "
         "left the sample. The lower unemployment of the households the cap misses fades over their four telephone "
         "waves as their circumstances change and as some of them would have left the panel, which is why the gap "
         "does not simply add up across the capped intakes."
         % (g[3], g[4], f2(-path_[3][1]), f2(-path_[4][1]), W.qname(W.qnext(W.LIMIT_TO, W.WAVES))), keep=True)
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(io.BytesIO(fig), width=Cm(16.3))
    pic.paragraph_format.space_after = Pt(2)
    pic.paragraph_format.keep_with_next = True
    para("Published unemployment rate, people aged 15 and over, not seasonally adjusted. Gap: published rate less the "
         "rate with six first-interview visits throughout.", size=8, color="56606A", after=10)
    para("Gap in the published rates, percentage points (published less six visits)", bold=True, size=9.5, after=3,
         keep=True)
    tb = doc.add_table(rows=6, cols=4)
    tb.style = "Table Grid"
    hdr = ["Quarter", "Capped intakes in the sample, by wave", "Unemployment rate", "Employment rate"]
    where = ["%d of 5, at wave%s %s" % (len(W.capped_waves(q)), "" if len(W.capped_waves(q)) == 1 else "s",
                                        " and ".join(", ".join(str(w) for w in W.capped_waves(q)).rsplit(", ", 1)))
             for q in AN.FORWARD]
    for j, t in enumerate(hdr):
        r = tb.cell(0, j).paragraphs[0].add_run(t)
        r.bold = True
        r.font.size = Pt(9)
    for i, q in enumerate(AN.FORWARD, start=1):
        vals = [short(q), where[i - 1], "+" + g[i - 1], e[i - 1]]
        for j, t in enumerate(vals):
            cell = tb.cell(i, j).paragraphs[0]
            if j >= 2:
                cell.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            rr = cell.add_run(t)
            rr.font.size = Pt(9)
    widths(tb, [2.8, 7.0, 3.3, 3.3])
    for row in tb.rows:
        for cell in row.cells:
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            cell.paragraphs[0].paragraph_format.keep_with_next = True
    para(after=4)
    para("How the gap is measured. The cap is simulated on %s, the one quarter whose five intakes were all first "
         "interviewed on six visits and all have their visit histories in the call-record file: for each quarter "
         "ahead, the households first interviewed at a fourth to sixth visit are taken out of the waves that "
         "quarter's capped intakes hold, the remaining sample is raked again to the population controls, and the gap "
         "is the change in each rate. Removing them from wave 1 alone understates every quarter and shows no gap at "
         "all once six visits return; carrying the wave-1 difference unchanged into the later waves overstates every "
         "quarter." % W.qname(W.TEMPLATE_Q), size=9)
    para("Regional field and flow figures for the quarter's quality report (first-interview visits to the July to "
         "September intake, the share of April to June's unemployed respondents in work by July to September, and "
         "total weekly hours actually worked) are in the accompanying workbook.", size=9)
    para("Raewyn Bailey", after=0)
    para(head_title, size=9, color="56606A")
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = footer.add_run("%s   page " % REF)
    fr.font.size = Pt(7.5)
    run = footer.add_run()
    run.font.size = Pt(7.5)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    doc.core_properties.title = "HLS visit cap note"
    doc.save(path)
    PP.normalize_ooxml(path, head_name, head_name, "2026-10-22T09:14:00Z", "2026-10-23T08:40:00Z",
                       title="HLS visit cap note")


def write_xlsx(path, c):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, Side
    h, a = c["h"], c["asks"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Gap path"
    ws["A1"] = "HLS: gap in the published rates from the three-visit cap, six visits from 1 April 2027"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Percentage points, published rate less the rate with six first-interview visits throughout. " \
               "People aged 15 and over. Simulated on %s, the one quarter whose five intakes were all " \
               "interviewed on six visits and all carry visit histories." % short(W.TEMPLATE_Q)
    ws["A2"].font = Font(italic=True, size=9, color="56606A")
    ws.append([])
    ws.append(["Quarter", "Unemployment rate gap", "Employment rate gap"])
    for q, g in zip(AN.FORWARD, h["path"]):
        ws.append([short(q), r2(g[0]), r2(g[1])])
    ws.append([])
    ws.append(["April to June to July to September 2026", "Unemployment rate", "Employment rate"])
    for lab, u, e in [("Published, April to June 2026", h["pub_q2"][0], h["pub_q2"][1]),
                      ("Published, July to September 2026", h["pub_q3"][0], h["pub_q3"][1]),
                      ("Gap, April to June 2026", h["gap_prior"][0], h["gap_prior"][1]),
                      ("Gap, July to September 2026", h["path"][0][0], h["path"][0][1]),
                      ("April to June 2026 on six visits throughout", h["without_q2_u"], h["without_q2_e"]),
                      ("July to September 2026 on six visits throughout", h["without_u"], h["without_e"]),
                      ("Change, published", h["rise_u"], h["rise_e"]),
                      ("of which the cap", h["cap_u"], h["cap_e"]),
                      ("of which the labour market", h["real_u"], h["real_e"])]:
        ws.append([lab, r2(u), r2(e)])
    thin = Side(style="thin", color="A3ABB3")
    for row in (4, 11):
        for cell in ws[row]:
            cell.font = Font(bold=True)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(wrap_text=True, vertical="bottom")
    for row in ws.iter_rows(min_row=5, max_row=20):
        for cell in row[1:]:
            if isinstance(cell.value, float):
                cell.number_format = "0.00" if cell.row in (12, 13, 16, 17) else "+0.00;-0.00;0.00"
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 22
    ws.freeze_panes = "A5"
    wr = wb.create_sheet("Regions")
    wr["A1"] = "HLS July to September 2026: field and flow figures by region"
    wr["A1"].font = Font(bold=True, size=12)
    wr.append([])
    wr.append(["Region", "First-interview visits, Jul-Sep 2026 intake", "Unemployed in Apr-Jun 2026",
               "of whom employed in Jul-Sep 2026", "Share employed (%)", "Total weekly hours actually worked (000)"])
    for g in W.RCODES:
        num, den = a["flows"]["golden_n"][g]
        wr.append([W.RNAME[g], a["visits"]["golden"][g], den, num, r2(a["flows"]["golden"][g]),
                   round(a["hours"]["golden"][g], 1)])
    for cell in wr[3]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="bottom")
        cell.border = Border(bottom=thin)
    for row in wr.iter_rows(min_row=4, max_row=9):
        row[1].number_format = "#,##0"
        row[4].number_format = "0.00"
        row[5].number_format = "#,##0.0"
    for col, wdt in zip("ABCDEF", (16, 18, 14, 14, 12, 18)):
        wr.column_dimensions[col].width = wdt
    wr.row_dimensions[3].height = 45
    wr.freeze_panes = "B4"
    wn = wb.create_sheet("Notes")
    for line in [
        "Visits: each visit to an address issued in the July to September 2026 intake, counted once per address and visit number.",
        "Flows: respondents unemployed in April to June 2026 who were interviewed again in July to September 2026, matched on address, sex and date of birth; unweighted.",
        "Hours: hours actually worked in the reference week by employed people, weighted with the final person weights; telephone interviews record hours and minutes.",
        "Gaps are rounded independently from unrounded estimates. The cap applies to the intakes of April to June 2026 to January to March 2027.",
        "Source: HLS extract of 20 October 2026; Merrowland Statistics, Labour Market Statistics.",
    ]:
        wn.append([line])
    wn.column_dimensions["A"].width = 130
    wb.save(path)
    PP.normalize_ooxml(path, W.PEOPLE["head"][0], W.PEOPLE["head"][0], "2026-10-22T08:31:00Z", "2026-10-23T08:36:00Z",
                       excel=True)


def main(target, golden, figures_path):
    os.makedirs(golden, exist_ok=True)
    c = compute(target)
    fig = chart_png(c)
    write_docx(os.path.join(golden, DOCX), c, fig)
    write_xlsx(os.path.join(golden, XLSX), c)
    with open(figures_path, "w") as fh:
        json.dump(figures(c), fh, indent=1, ensure_ascii=False)
        fh.write("\n")
    return c
