"""Deterministic writers (adapted from the house pattern). Every OOXML file is repacked with fixed entry times after the producer scrub;
every PDF is written with ReportLab's invariant mode and then scrubbed at equal byte length; parquet,
csv, sqlite, txt and eml are written byte-stable. File times are set to one in-fiction date at the end."""
import csv
import datetime as dt
import importlib.util
import io
import os
import sqlite3
import zipfile
from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent
SCRUB_PATH = REPO / ".claude/skills/reduce-house-fixes/scripts/scrub_producer_metadata.py"
_spec = importlib.util.spec_from_file_location("scrub_producer_metadata", SCRUB_PATH)
SCRUB = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(SCRUB)

FLOOR, CEILING = "2023-01-01", "2027-03-03"


def write_csv(path, rows, columns):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(columns)
        for r in rows:
            w.writerow(["" if r.get(c) is None else r.get(c) for c in columns])
    return len(rows)


def write_parquet(path, columns, types, rows, row_group=100_000):
    arrays = []
    for c, t in zip(columns, types):
        vals = [r[c] for r in rows]
        if t == "ts":
            arr = pa.array([None if v in ("", None) else dt.datetime.strptime(v, "%Y-%m-%d %H:%M") for v in vals],
                           pa.timestamp("s"))
        elif t == "date":
            arr = pa.array([None if v in ("", None) else dt.date.fromisoformat(v) for v in vals], pa.date32())
        elif t == "int":
            arr = pa.array([None if v in ("", None) else int(v) for v in vals], pa.int32())
        else:
            arr = pa.array([None if v in ("", None) else str(v) for v in vals], pa.string())
        arrays.append(arr)
    table = pa.Table.from_arrays(arrays, names=columns).replace_schema_metadata(None)
    pq.write_table(table, str(path), compression="snappy", row_group_size=row_group, write_statistics=True,
                   use_dictionary=True, store_schema=False)
    return len(rows)


def repack_ooxml(path, when):
    """Rewrite the archive with the same entries in the same order and a fixed entry time."""
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        data = [(i.filename, z.read(i.filename), i.compress_type) for i in infos]
    tmp = str(path) + ".tmp"
    with zipfile.ZipFile(tmp, "w") as out:
        for name, blob, ct in data:
            zi = zipfile.ZipInfo(name, date_time=when.timetuple()[:6])
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            zi.create_system = 0
            out.writestr(zi, blob)
    os.replace(tmp, path)


def scrub_ooxml(path, producer, when):
    stamp = when.strftime("%Y-%m-%d")
    SCRUB.scrub_ooxml(str(path), producer, stamp, FLOOR, CEILING)
    repack_ooxml(path, when)


def scrub_pdf(path, producer, when):
    stamp = when.strftime("%Y-%m-%d")
    SCRUB.scrub_pdf(str(path), producer, stamp, FLOOR, CEILING)


# ---------------------------------------------------------------------------------------- docx
def write_docx(path, title, blocks, author, when, producer, org="Wealdmoor County Council"):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2.0)
    hdr = doc.sections[0].header.paragraphs[0]
    hdr.text = org
    hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(14)
    for kind, val in blocks:
        if kind == "meta":
            q = doc.add_paragraph()
            r = q.add_run(val)
            r.italic = True
            r.font.size = Pt(9.5)
        elif kind == "h":
            q = doc.add_paragraph()
            r = q.add_run(val)
            r.bold = True
            r.font.size = Pt(11)
        elif kind == "p":
            doc.add_paragraph(val)
        elif kind == "table":
            t = doc.add_table(rows=len(val), cols=len(val[0]))
            t.style = "Table Grid"
            for i, row in enumerate(val):
                for j, cell in enumerate(row):
                    t.cell(i, j).text = cell
                    if i == 0:
                        for rr in t.cell(i, j).paragraphs[0].runs:
                            rr.bold = True
        elif kind == "foot":
            ft = doc.sections[0].footer.paragraphs[0]
            ft.text = val
            for rr in ft.runs:
                rr.font.size = Pt(8)
    cp = doc.core_properties
    cp.author = author
    cp.last_modified_by = author
    cp.title = title
    cp.comments = ""
    cp.subject = ""
    cp.keywords = ""
    cp.category = ""
    cp.revision = 4
    cp.created = when - dt.timedelta(days=9, hours=3)
    cp.modified = when
    doc.save(str(path))
    scrub_ooxml(path, producer, when)


# ---------------------------------------------------------------------------------------- pdf
def write_pdf(path, title, sub, by, blocks, author, when, producer, fields=None):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=12.5, spaceAfter=5)
    small = ParagraphStyle("s", parent=body, fontSize=8.5, leading=10.5)
    h1 = ParagraphStyle("h1", parent=body, fontName="Helvetica-Bold", fontSize=13, leading=16, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=body, fontName="Helvetica-Bold", fontSize=10.5, leading=13, spaceBefore=6)
    meta = ParagraphStyle("m", parent=body, fontName="Helvetica-Oblique", fontSize=9, leading=11)
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=2.1 * cm, rightMargin=2.1 * cm, topMargin=1.8 * cm,
                            bottomMargin=1.8 * cm, title=title, author=author, subject="", creator=author,
                            invariant=1)
    story = [Paragraph(escape(title), h1)]
    if sub:
        story.append(Paragraph(escape(sub), meta))
    if by:
        story.append(Paragraph(escape(by), meta))
    story.append(Spacer(1, 6))
    for kind, val in blocks:
        if kind == "h":
            story.append(Paragraph(escape(val), h2))
        elif kind == "p":
            story.append(Paragraph(escape(val), body))
        elif kind == "small":
            story.append(Paragraph(escape(val), small))
        elif kind == "table":
            data = [[Paragraph(escape(str(c)), small) for c in row] for row in val["rows"]]
            t = Table(data, colWidths=val.get("widths"), repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.9, 0.9, 0.9)),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story.append(t)
            story.append(Spacer(1, 6))
    doc.build(story)
    scrub_pdf(path, producer, when)


# ---------------------------------------------------------------------------------------- xlsx
def xlsx_book(path, title, author, company, when):
    import xlsxwriter
    wb = xlsxwriter.Workbook(str(path), {"strings_to_numbers": False})
    wb.set_properties({"title": title, "author": author, "company": company, "created": when})
    return wb


def close_xlsx(wb, path, producer, when):
    wb.close()
    scrub_ooxml(path, producer, when)


# ---------------------------------------------------------------------------------------- sqlite
def write_sqlite(path, schema, tables):
    if os.path.exists(path):
        os.remove(path)
    con = sqlite3.connect(str(path))
    con.execute("PRAGMA page_size=4096")
    con.execute("PRAGMA journal_mode=OFF")
    for ddl in schema:
        con.execute(ddl)
    for name, cols, rows in tables:
        con.executemany("INSERT INTO %s (%s) VALUES (%s)" % (name, ",".join(cols), ",".join("?" * len(cols))), rows)
    con.commit()
    con.execute("VACUUM")
    con.close()


def set_mtimes(root, when):
    ts = when.timestamp()
    for p in sorted(Path(root).rglob("*")):
        if p.is_file():
            os.utime(p, (ts, ts))


def write_text(path, text):
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(text)


def write_json(path, obj):
    import json
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(json.dumps(obj, indent=2, ensure_ascii=False) + "\n")


def set_mtime(path, when):
    ts = when.timestamp()
    os.utime(path, (ts, ts))
