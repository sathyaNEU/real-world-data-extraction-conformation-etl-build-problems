"""The analysis on the shipped files alone: load target/, rebuild each quarter's sample, rake to the controls file,
and compute the published rates, the visit cap's gap path (the golden) and each wrong path the ladder names."""
import os

import numpy as np
import openpyxl
import pandas as pd

import records as R
import world as W

AGE_EDGES = np.array([25, 35, 45, 55, 65, 75])


def load(target):
    p = pd.read_parquet(os.path.join(target, R.PERSON))
    p["serial9"] = p["serial"].str.slice(0, 9)
    c = pd.read_csv(os.path.join(target, R.CALLS), dtype={"serial": str, "area_code": str})
    c["serial9"] = c["serial"].str.slice(0, 9)
    interview = c[c["outcome"] == 110].groupby("serial9")["visit_no"].min()
    p["interview_visit"] = p["serial9"].map(interview)
    wb = openpyxl.load_workbook(os.path.join(target, R.CONTROLS), data_only=True)
    ctl = {}
    reg = {r[0]: r for r in wb["Region"].iter_rows(min_row=2, values_only=True)}
    agesex = {r[0]: r for r in wb["Age and sex"].iter_rows(min_row=2, values_only=True)}
    for q in reg:
        ctl[q] = (np.array(reg[q][2:8], float), np.array(agesex[q][2:16], float))
    return dict(person=p, calls=c, controls=ctl)


def cells(df):
    g = np.searchsorted(AGE_EDGES, df["age"].to_numpy(), side="right")
    return 2 * g + (df["sex"].to_numpy() - 1)


def rake(df, controls, iters=200):
    rt, ct = controls
    region = df["region"].map({c: i for i, c in enumerate(W.RCODES)}).to_numpy()
    cell = cells(df)
    w = np.ones(len(df))
    for _ in range(iters):
        w *= (rt / np.bincount(region, w, 6))[region]
        f = ct / np.bincount(cell, w, 14)
        w *= f[cell]
        if np.max(np.abs(f - 1)) < 1e-12:
            break
    return w


def rates(df, w):
    s = df["ilo_status"].to_numpy()
    e, u = w[s == 1].sum(), w[s == 2].sum()
    return 100 * u / (e + u), 100 * e / w.sum()


def quarter(d, q):
    return d["person"][d["person"]["quarter"] == W.qlabel(q)].reset_index(drop=True)


def published(d, q, reweight="file"):
    df = quarter(d, q)
    w = df["weight"].to_numpy() if reweight == "file" else rake(df, d["controls"][W.qlabel(q)])
    return rates(df, w)


def positions(k):
    return tuple(x for x in (1 + k, k) if 1 <= x <= W.WAVES)


FORWARD = [W.qnext(W.LIMIT_FROM, k) for k in range(5)]


def gap_path(d, template=W.TEMPLATE_Q, reweight="rake", carry="panel"):
    """The gap (published less six-visit) in each quarter July to September 2026 to July to September 2027.
    carry: 'panel' (golden: the capped intakes' late households dropped at every wave they occupy) or 'wave1'
    (the default: dropped at wave 1 while the cap runs, nothing from January 2027).
    reweight: 'rake' (golden) or 'keep' (the template's final weights kept)."""
    df = quarter(d, template)
    ctl = d["controls"][W.qlabel(template)]
    w0 = rake(df, ctl)
    base = rates(df, w0)
    late = (df["interview_visit"] > W.CAP_VISITS).to_numpy()
    wave = df["wave"].to_numpy()
    out = []
    for k, q in enumerate(FORWARD):
        pos = positions(k) if carry == "panel" else ((1,) if W.capped(q) else ())
        keep = ~(np.isin(wave, pos) & late)
        sub = df[keep].reset_index(drop=True)
        w = rake(sub, ctl) if reweight == "rake" else w0[keep]
        r = rates(sub, w)
        out.append((r[0] - base[0], r[1] - base[1]))
    return out


def constant_gap_path(d, template=W.TEMPLATE_Q):
    """Rung 1: the wave-1 gap (the wave re-estimated alone without its late households) carried at each wave a
    capped intake occupies, weighted by that wave's share of the labour force (unemployment rate) or population."""
    df = quarter(d, template)
    w0 = rake(df, d["controls"][W.qlabel(template)])
    s = df["ilo_status"].to_numpy()
    wave = df["wave"].to_numpy()
    late = (df["interview_visit"] > W.CAP_VISITS).to_numpy()
    act = s <= 2
    w1 = wave == 1
    full = rates(df[w1], w0[w1])
    cut = rates(df[w1 & ~late], w0[w1 & ~late])
    du, de = cut[0] - full[0], cut[1] - full[1]
    sa = [(w0 * act)[wave == v].sum() / (w0 * act).sum() for v in range(1, 6)]
    sp = [w0[wave == v].sum() / w0.sum() for v in range(1, 6)]
    return [(sum(du * sa[p - 1] for p in positions(k)), sum(de * sp[p - 1] for p in positions(k))) for k in range(5)]


def headline(d):
    """The call's figures: published rates in April to June and July to September 2026, the cap's July to September
    gap, the rate without the cap and the labour market's share of the change."""
    q2, q3 = published(d, (2026, 2)), published(d, (2026, 3))
    path = gap_path(d)
    g3u, g3e = path[0]
    return dict(path=path, pub_q2=q2, pub_q3=q3, rise_u=q3[0] - q2[0], rise_e=q3[1] - q2[1],
                without_u=q3[0] - g3u, without_e=q3[1] - g3e, real_u=q3[0] - g3u - q2[0], real_e=q3[1] - g3e - q2[1],
                series=[(q, published(d, q)) for q in W.qrange(W.FIRST_Q, W.LAST_Q)])
