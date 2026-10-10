"""Golden deliverables for task125:  python3 task125/generator/golden.py

Reads only task125/target/. Rebuilds the digit screen from the published payments, forecasts the
payments the filter routes in 2027/28, reads the 2025/26 run log, acknowledgements and call-off papers,
and writes the two files the prompt names into task125/golden/:
  examination_calloff_2027-28.docx  (the order note, with the monthly chart rendered here)
  examination_calloff_2027-28.xlsx  (the plan-year months, 2025/26 by month, 2025/26 by quarter)
Prints the critical components' figures at the end."""
import csv
import datetime as dt
import io
import itertools
import json
import math
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from pypdf import PdfReader

HERE = Path(__file__).resolve().parent
TASK = HERE.parent
T = TASK / "target"
OUT = TASK / "golden"
sys.path.insert(0, str(HERE))
import writers as W  # noqa: E402  (deterministic save and producer scrub only; no data)

BEN = {d: math.log10(1 + 1 / d) for d in range(10, 100)}
YEARS = ["2023/24", "2024/25", "2025/26"]
PLAN0, PLAN1 = dt.date(2027, 4, 1), dt.date(2028, 3, 31)
SHORT = {"Adult Social Care": "ASC", "Housing Support": "HS", "Highways & Transport": "HT",
         "Property & Facilities": "PF", "Waste & Environment": "WE"}


def find(pattern):
    hits = sorted(T.glob(pattern))
    assert len(hits) == 1, (pattern, hits)
    return hits[0]


def pdf_text(p):
    return "\n".join(pg.extract_text() for pg in PdfReader(str(p)).pages)


def pence(s):
    neg = s.startswith("-")
    a, b = s.lstrip("-").split(".")
    v = int(a) * 100 + int(b)
    return -v if neg else v


def fy(d):
    y = d.year if d.month >= 4 else d.year - 1
    return "%d/%02d" % (y, (y + 1) % 100)


def cell(v):
    return int(str(v // 100)[:2])


def ym(d):
    return "%d-%02d" % (d.year, d.month)


def month_list(y0, m0, n):
    out = []
    for i in range(n):
        t = y0 * 12 + m0 - 1 + i
        out.append("%d-%02d" % (t // 12, t % 12 + 1))
    return out


# ---------------------------------------------------------------------------------------- the payments
PAYS = []   # (dept, date, gross pence, expense_type, vendor_no)
NETVAT = []
with open(find("wealdmoor_spend_over_500_*.csv"), encoding="utf-8") as fh:
    for r in csv.DictReader(fh):
        d = dt.datetime.strptime(r["payment_date"], "%d/%m/%Y").date()
        n, v = pence(r["net_amount"]), pence(r["vat_amount"])
        PAYS.append((r["department"], d, n + v, r["expense_type"], r["vendor_no"]))
        NETVAT.append((n, v))
DEPTS = list(dict.fromkeys(p[0] for p in PAYS))
LAST_PAY = max(p[1] for p in PAYS)

# ---------------------------------------------------------------------------------------- the screen
st_text = pdf_text(find("digit_conformity_statements_*.pdf"))
STATEMENTS = {}
for block in re.split(r"Digit-conformity statement ", st_text)[1:]:
    year = block[:7]
    for d in DEPTS:
        m = re.search(re.escape(d) + r"\s*\n\s*([\d,]+)\s*\n\s*(0\.\d{5})", block)
        STATEMENTS[(year, d, "tested")] = int(m.group(1).replace(",", ""))
        STATEMENTS[(year, d, "mad")] = m.group(2)
    STATEMENTS[(year, "ALL", "mad")] = re.search(r"Council \(all departments\)\s*\n\s*(0\.\d{5})", block).group(1)
assert len(STATEMENTS) == 63


def in_range(g):
    """Payments as made (gross), complete decades 1,000.00 to 999,999.99; credit notes excluded."""
    return 100000 <= g <= 99999999


SCREEN = defaultdict(Counter)
for dept, d, g, _, _ in PAYS:
    if fy(d) in YEARS and in_range(g):
        SCREEN[(fy(d), dept)][cell(g)] += 1


def mad(c):
    n = sum(c.values())
    return sum(abs(c.get(k, 0) / n - BEN[k]) for k in range(10, 100)) / 90


hits = 0
for y in YEARS:
    pooled = Counter()
    for d in DEPTS:
        pooled.update(SCREEN[(y, d)])
        hits += sum(SCREEN[(y, d)].values()) == STATEMENTS[(y, d, "tested")]
        hits += "%.5f" % mad(SCREEN[(y, d)]) == STATEMENTS[(y, d, "mad")]
    hits += "%.5f" % mad(pooled) == STATEMENTS[(y, "ALL", "mad")]
assert hits == 63, hits
REPRODUCED = hits


def flagged(y):
    out = set()
    for d in DEPTS:
        c = SCREEN[(y, d)]
        n = sum(c.values())
        for k in range(10, 100):
            e = BEN[k] * n
            if c.get(k, 0) > 1.5 * e and c.get(k, 0) - e >= 30:
                out.add((d, k))
    return out


CELLS = {y: flagged(y) for y in YEARS}
PLAN_CELLS = CELLS["2025/26"]        # methodology 5: latest statement published when the order is placed
assert len(PLAN_CELLS) == 5

# ---------------------------------------------------------------------------------------- BACS calendar
CAL = openpyxl.load_workbook(find("bacs_payment_calendar_*.xlsx"), read_only=True)
SUB = {}                     # payment date -> BACS submission date
CAL_ROWS = defaultdict(list)  # sheet -> (type, payment date, submission date)
for sh in CAL.worksheets:
    for r in sh.iter_rows(min_row=4, values_only=True):
        if r and r[0]:
            pd_ = dt.datetime.strptime(r[1], "%d/%m/%Y").date()
            sd = dt.datetime.strptime(r[2], "%d/%m/%Y").date()
            assert SUB.get(pd_, sd) == sd
            SUB[pd_] = sd
            CAL_ROWS[sh.title].append((r[0], pd_, sd))


def run_month(d):
    """The filter's run month: the month the payment's BACS file is submitted (run book section 2)."""
    return ym(SUB[d])


# ---------------------------------------------------------------------------------------- 2025/26 streams in the plan cells
SL_EXP, TSP_EXP = "Shared Lives carer payments", "Tenancy sustainment payments"
FLAT = Counter()             # (dept, cell) -> 2025/26 count, every stream except carers and instalments
FLAT_MONTH = defaultdict(Counter)
for dept, d, g, exp, ven in PAYS:
    if fy(d) == "2025/26" and in_range(g) and (dept, cell(g)) in PLAN_CELLS and exp not in (SL_EXP, TSP_EXP):
        FLAT[(dept, cell(g))] += 1
        FLAT_MONTH[(dept, cell(g))][run_month(d)] += 1
for k, c in FLAT_MONTH.items():      # every one of these streams is level: the same count in every run month
    assert len(c) == 12 and len(set(c.values())) == 1, (k, c)
FLAT_PER_MONTH = {k: next(iter(set(c.values()))) for k, c in FLAT_MONTH.items()}
FLAT_YEAR = sum(FLAT.values())

# ---------------------------------------------------------------------------------------- Housing Support: the instalment term
TSP = defaultdict(list)
ROUND = {}
with open(find("tenancy_sustainment_payments_*.csv"), encoding="utf-8") as fh:
    for r in csv.DictReader(fh):
        TSP[r["case_ref"]].append((dt.date.fromisoformat(r["payment_date"]), pence(r["amount"])))
        ROUND[r["case_ref"]] = r["scheme_round"]
done = [len(v) for k, v in TSP.items() if ROUND[k] == "2021"]
TERMS_SEEN = Counter(done)
assert len(done) == 288 and set(done) == {30}
TERM = 30
assert all(cell(a) == 14 for v in TSP.values() for _, a in v)
TSP_SCHED = Counter()        # month -> instalments due, every household run out on the 30-instalment term
for k, v in TSP.items():
    first = min(x[0] for x in v)
    for m in month_list(first.year, first.month, TERM):
        TSP_SCHED[m] += 1
PLAN_MONTHS = month_list(2027, 4, 12)
HS_INSTAL = sum(TSP_SCHED[m] for m in PLAN_MONTHS)
LIVE_2425 = sum(1 for k in TSP if ROUND[k] != "2021")
# the scheme pays on the 15th: each instalment's BACS file is submitted in the month it is paid
for typ, pd_, sd in CAL_ROWS["2027-28"]:
    if typ == "Tenancy sustainment":
        assert ym(pd_) == ym(sd)

# ---------------------------------------------------------------------------------------- Shared Lives: what each carer is paid after 31 March 2027
rt = pdf_text(find("shared_lives_carer_rates_*.pdf"))
RATES = {m.group(1): int(m.group(2)) for m in
         re.finditer(r"(Long-term, band \d|Home First step-down, \w+)\s*\n?\s*(\d{3})\.00", rt)}
assert len(RATES) == 6, RATES
CARER_H = defaultdict(dict)
for dept, d, g, exp, ven in PAYS:
    if exp == SL_EXP:
        assert d not in CARER_H[ven]
        CARER_H[ven][d] = g
SL_DATES = sorted({d for h in CARER_H.values() for d in h})
CARER = {v: h[SL_DATES[-1]] for v, h in CARER_H.items() if SL_DATES[-1] in h}     # paid on the latest run
DECOMP = defaultdict(list)
for n in (1, 2, 3):
    for ms in itertools.combinations_with_replacement(sorted(RATES), n):
        DECOMP[400 * sum(RATES[x] for x in ms)].append(ms)
# a carer's amount is four weeks of one set of placements, uniquely; the amounts that are not are each half of
# a household's fee, paid to two consecutive vendor numbers on the same runs (a household approved jointly)
HALF = {v for v, a in CARER.items() if not DECOMP[a]}
assert all(len(DECOMP[a]) == 1 for v, a in CARER.items() if v not in HALF)
assert all(len(DECOMP[2 * CARER[v]]) == 1 for v in HALF)
_hv = sorted(HALF)
assert len(_hv) % 2 == 0 and all(int(_hv[i + 1]) == int(_hv[i]) + 1 and CARER_H[_hv[i]] == CARER_H[_hv[i + 1]]
                                 for i in range(0, len(_hv), 2))
JOINT_HH = len(_hv) // 2
# the joint households that lost a guest inside the extract: how a household is paid with one guest left
CHANGES = []
for v in sorted(CARER_H):
    w = "%06d" % (int(v) + 1)
    if w not in CARER_H:
        continue
    both = sorted(set(CARER_H[v]) & set(CARER_H[w]))
    h0 = CARER_H[v][both[0]] if both else None
    if not both or h0 != CARER_H[w][both[0]] or DECOMP[h0] or len(DECOMP[2 * h0]) != 1:
        continue
    moved = [d for d in SL_DATES if d in CARER_H[v] and CARER_H[v][d] != h0]
    if not moved:
        continue
    d0 = moved[0]
    if d0 in CARER_H[w]:
        CHANGES.append((d0, "halves", len(DECOMP[2 * CARER_H[v][d0]][0])))
    else:
        assert not any(d >= d0 for d in CARER_H[w])
        CHANGES.append((d0, "whole", len(DECOMP[CARER_H[v][d0]][0])))
assert all((k == "whole") == (n == 1) for d, k, n in CHANGES) and {k for d, k, n in CHANGES} == {"whole", "halves"}
ONE_GUEST_DATES = sorted(d for d, k, n in CHANGES if k == "whole")
AFTER = {}
for ven, a in CARER.items():
    if ven not in HALF:
        AFTER[ven] = 400 * sum(RATES[x] for x in DECOMP[a][0] if not x.startswith("Home First"))
for i in range(0, len(_hv), 2):
    left = [x for x in DECOMP[2 * CARER[_hv[i]]][0] if not x.startswith("Home First")]
    fee = 400 * sum(RATES[x] for x in left)
    AFTER[_hv[i]], AFTER[_hv[i + 1]] = (fee, 0) if len(left) == 1 else (fee // 2, fee // 2)


def _in14(a):
    return bool(a) and a >= 100000 and cell(a) == 14


SL14_NOW = sum(1 for a in CARER.values() if cell(a) == 14)
SL14_AFTER = sum(1 for a in AFTER.values() if _in14(a))
SL_MOVED = sum(1 for v in CARER if v not in HALF and cell(CARER[v]) != 14 and _in14(AFTER[v]))
HALF_IN, WHOLE_IN = [], []
for i in range(0, len(_hv), 2):
    va, vb = _hv[i], _hv[i + 1]
    if AFTER[vb] and _in14(AFTER[va]):
        HALF_IN += [va, vb]                   # still paid in halves, and the halves enter cell 14
    elif not AFTER[vb] and _in14(AFTER[va]):
        WHOLE_IN.append(va)                   # one guest left: the whole fee to the first vendor, in cell 14
HALF_MOVED = len(HALF_IN)                     # halves that stay halves and enter cell 14
WHOLE_MOVED = len(WHOLE_IN)                   # one-guest households paid whole into cell 14
HALF_HH, WHOLE_HH = HALF_MOVED // 2, WHOLE_MOVED
HALF_AMTS = sorted({CARER[v] for v in HALF_IN})
HALF_AFTER = sorted({AFTER[v] for v in HALF_IN})
WHOLE_AMTS = sorted({CARER[v] for v in WHOLE_IN})
WHOLE_AFTER = sorted({AFTER[v] for v in WHOLE_IN})
WHOLE_HALVED = sorted({a // 2 for a in WHOLE_AFTER})
assert SL_MOVED + HALF_MOVED + WHOLE_MOVED + SL14_NOW == SL14_AFTER
SL_RUNS = [(pd_, sd) for typ, pd_, sd in CAL_ROWS["2027-28"] if typ == "Shared Lives carers"]
assert len(SL_RUNS) == 13 and all(PLAN0 <= p <= PLAN1 for p, _ in SL_RUNS)

# ---------------------------------------------------------------------------------------- the 2027/28 forecast
ASC_OTHER = FLAT[("Adult Social Care", 14)]
HS_OTHER = FLAT[("Housing Support", 14)]
BY_CELL = dict(FLAT)
BY_CELL[("Adult Social Care", 14)] = ASC_OTHER + len(SL_RUNS) * SL14_AFTER
BY_CELL[("Housing Support", 14)] = HS_OTHER + HS_INSTAL
ORDER_U = sum(BY_CELL.values())
ORDER = int(round(ORDER_U, -2))
BY_DEPT = Counter()
for (d, c), n in BY_CELL.items():
    BY_DEPT[d] += n
TOP_DEPT, TOP_N = max(BY_DEPT.items(), key=lambda x: x[1])
RUNG3 = ORDER_U - len(SL_RUNS) * (SL14_AFTER - SL14_NOW)
RUNG5 = ORDER_U - len(SL_RUNS) * WHOLE_MOVED
RUNG4 = ORDER_U - len(SL_RUNS) * (HALF_MOVED + WHOLE_MOVED)

# by run month (ask A): the month each payment's BACS file is submitted
PLAN_BY = {m: Counter() for m in PLAN_MONTHS}
for (d, c), per in FLAT_PER_MONTH.items():
    for m in PLAN_MONTHS:
        PLAN_BY[m][d] += per
for m in PLAN_MONTHS:
    PLAN_BY[m]["Housing Support"] += TSP_SCHED[m]
for p, s in SL_RUNS:
    PLAN_BY[ym(s)]["Adult Social Care"] += SL14_AFTER
A = [sum(PLAN_BY[m].values()) for m in PLAN_MONTHS]
assert sum(A) == ORDER_U, (sum(A), ORDER_U)
BUSY = PLAN_MONTHS[A.index(max(A))]
assert A.count(max(A)) == 1

# ---------------------------------------------------------------------------------------- the 2025/26 run log (Fernhollow's figure, ask B1)
CODE = {}
for d in DEPTS:
    CODE["".join(w[0] for w in re.findall(r"[A-Za-z']+", d) if w[0].isupper())] = d
rl = openpyxl.load_workbook(find("digit_filter_run_log_*.xlsx"), read_only=True).worksheets[0]
hdr, LOG = None, []
for r in rl.iter_rows(values_only=True):
    if r and r[0] == "run_ref":
        hdr = list(r)
    elif hdr and r and r[0]:
        LOG.append(dict(zip(hdr, r)))
for r in LOG:
    r["dept"] = CODE.get(r["department"]) or next(d for d in DEPTS if d.upper().startswith(r["department"][:2]))
LOG_CELLS = {(r["dept"], int(r["flagged_cell"])) for r in LOG if r["flagged_cell"] not in (None, "")}
assert LOG_CELLS == CELLS["2023/24"]
RERUN = {r["run_month"] for r in LOG if r["run_type"] == "Re-run"}
B1_BY = defaultdict(Counter)          # run month -> dept -> routed; a re-run replaces, a supplementary run adds
for r in LOG:
    if r["run_type"] == "Scheduled" and r["run_month"] in RERUN:
        continue
    B1_BY[r["run_month"]][r["dept"]] += int(r["payments_routed"])
M2526 = month_list(2025, 4, 12)
B1 = [sum(B1_BY[m].values()) for m in M2526]
FERN = sum(B1)
FERN_H = int(round(FERN, -2))
GAP = FERN_H - ORDER
assert int(round(FERN - ORDER_U, -2)) == GAP
# the log ties to the spending file by BACS submission month, on the cells it applied
tie = Counter()
for dept, d, g, _, _ in PAYS:
    if in_range(g) and d in SUB and ym(SUB[d]) in M2526 and (dept, cell(g)) in LOG_CELLS:
        tie[ym(SUB[d])] += 1
assert [tie[m] for m in M2526] == B1

# ---------------------------------------------------------------------------------------- examinations (ask B2) and charges (ask C)
BATCH_MONTH = {r["batch_ref"]: r["run_month"] for r in LOG if r["batch_ref"]}
ACKS = json.loads(find("fernhollow_batch_acknowledgements_*.json").read_text(encoding="utf-8"))["batches"]
B2c = Counter()
for a in ACKS:
    B2c[BATCH_MONTH[a["batch_ref"]]] += a["payments_examined"]
B2 = [B2c[m] for m in M2526]
q2 = pdf_text(find("fernhollow_service_report_q2_*.pdf"))
assert int(re.search(r"Payments examined\s*\n?\s*([\d,]+)", q2).group(1).replace(",", "")) == sum(B2[3:6])

terms = pdf_text(find("pasf_lot2_calloff_terms_*.pdf"))
MIN_BATCH = int(re.search(r"minimum charge of (\d+) examinations", terms).group(1))
SHARE = int(re.search(r"at (\d+) per cent of the base rate", terms).group(1)) / 100
od = pdf_text(find("calloff_*order_and_variation*.pdf"))
ALLOC0 = int(re.search(r"Quarterly allocation\s*\n?\s*([\d,]+)", od).group(1).replace(",", ""))
ALLOC1 = int(re.search(r"increased to\s*\n?\s*([\d,]+)", od).group(1).replace(",", ""))
ALLOC = {1: ALLOC0, 2: ALLOC0, 3: ALLOC1, 4: ALLOC1}      # the call-off as varied from 1 October 2025
card = {r[0]: r[1] for r in openpyxl.load_workbook(find("fernhollow_rate_cards_*.xlsx"), read_only=True)
        ["2025-26"].iter_rows(values_only=True) if r[0]}
BASE = float(next(v for k, v in card.items() if str(k).startswith("Base rate")))
QOF = {4: 1, 5: 1, 6: 1, 7: 2, 8: 2, 9: 2, 10: 3, 11: 3, 12: 3, 1: 4, 2: 4, 3: 4}
CHARGED = Counter()
for a in ACKS:
    ex = a["payments_examined"]
    CHARGED[QOF[int(BATCH_MONTH[a["batch_ref"]][5:])]] += max(ex, MIN_BATCH) if ex else 0
PREM = [max(0, CHARGED[q] - ALLOC[q]) for q in (1, 2, 3, 4)]
UNUSED = [max(0, ALLOC[q] - CHARGED[q]) for q in (1, 2, 3, 4)]
UNUSED_GBP = [int(round(u * SHARE * BASE)) for u in UNUSED]
assert all(abs(u * SHARE * BASE * 100 % 100 - 50) > 1e-6 for u in UNUSED)   # no half-pound tie
# invoice quantities net of the credit note equal the charged examinations
inv_q = Counter()
for r in csv.DictReader(open(find("fernhollow_invoice_lines_*.csv"), encoding="utf-8")):
    inv_q[r["batch_ref"]] += int(r["quantity"])
assert all(inv_q[a["batch_ref"]] == (max(a["payments_examined"], MIN_BATCH) if a["payments_examined"] else 0)
           for a in ACKS)

# ---------------------------------------------------------------------------------------- the chart's history, April 2025 to March 2027
# 2025/26: the run log. 2026/27: the 2024/25 screen's cells (the statement published when that call-off was
# placed), counted from the spending file by run month to January 2027; February and March 2027 run past
# the extract (payments to 25 February 2027) and are projected on the same streams.
CELLS_2627 = CELLS["2024/25"]
M2627 = month_list(2026, 4, 12)
OBS_2627 = [m for m in M2627 if m <= "2027-01"]
HIST = {m: Counter(B1_BY[m]) for m in M2526}
for m in M2627:
    HIST[m] = Counter()
for dept, d, g, _, _ in PAYS:
    if in_range(g) and d in SUB and ym(SUB[d]) in OBS_2627 and (dept, cell(g)) in CELLS_2627:
        HIST[ym(SUB[d])][dept] += 1
extra = CELLS_2627 - PLAN_CELLS
we_obs = [sum(1 for dept, d, g, _, _ in PAYS if in_range(g) and d in SUB and ym(SUB[d]) == m
              and (dept, cell(g)) in extra) for m in OBS_2627]
WE_PROJ = int(round(sum(we_obs) / len(we_obs)))
PROJECTED = [m for m in M2627 if m not in OBS_2627]
for m in PROJECTED:
    for (d, c), per in FLAT_PER_MONTH.items():
        HIST[m][d] += per
    HIST[m]["Housing Support"] += TSP_SCHED[m]
    for (d, c) in extra:
        HIST[m][d] += WE_PROJ
for typ, p, s in CAL_ROWS["2026-27"]:
    if typ == "Shared Lives carers" and ym(s) in PROJECTED:
        HIST[ym(s)]["Adult Social Care"] += SL14_NOW
for m in PLAN_MONTHS:
    HIST[m] = PLAN_BY[m]
CHART_MONTHS = M2526 + M2627 + PLAN_MONTHS

# ---------------------------------------------------------------------------------------- the bridge from Fernhollow's figure to the order
Y2526 = Counter()            # 2025/26 payments in the plan cells, by payment date
for dept, d, g, _, _ in PAYS:
    if fy(d) == "2025/26" and in_range(g) and (dept, cell(g)) in PLAN_CELLS:
        Y2526[(dept, cell(g))] += 1
WE_2526 = FERN - sum(Y2526.values())
assert WE_2526 == sum(B1_BY[m][d] for m in M2526 for (d, c) in LOG_CELLS - PLAN_CELLS)
HS_DELTA = BY_CELL[("Housing Support", 14)] - Y2526[("Housing Support", 14)]
SL_DELTA = BY_CELL[("Adult Social Care", 14)] - Y2526[("Adult Social Care", 14)]
assert FERN - WE_2526 + HS_DELTA + SL_DELTA == ORDER_U
assert SL_DELTA == len(SL_RUNS) * (SL_MOVED + HALF_MOVED + WHOLE_MOVED)
SL_DELTA_SINGLE, SL_DELTA_JOINT = len(SL_RUNS) * SL_MOVED, len(SL_RUNS) * HALF_MOVED
SL_DELTA_WHOLE = len(SL_RUNS) * WHOLE_MOVED
HS_INST_2526 = Y2526[("Housing Support", 14)] - HS_OTHER
DUAL_NOW = sorted({CARER[v] // 100 for v in CARER if v not in HALF and cell(CARER[v]) != 14 and AFTER[v]
                   and cell(AFTER[v]) == 14})
DUAL_AFTER = sorted({AFTER[v] // 100 for v in CARER if v not in HALF and cell(CARER[v]) != 14 and AFTER[v]
                     and cell(AFTER[v]) == 14})
assert len(DUAL_AFTER) == 1 and len(HALF_AFTER) == 1 and len(WHOLE_AFTER) == 1


def fmt(n):
    return "{:,}".format(n)


MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def mlabel(m, long=False):
    y, mm = int(m[:4]), int(m[5:])
    return ("%s %d" % (["January", "February", "March", "April", "May", "June", "July", "August", "September",
                         "October", "November", "December"][mm - 1], y)) if long else "%s %d" % (MON[mm - 1], y)


# ---------------------------------------------------------------------------------------- the chart
STACK = ["Adult Social Care", "Housing Support", "Highways & Transport", "Property & Facilities",
         "Waste & Environment"]
COLOURS = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4"]   # dataviz reference slots 1 to 5, validated


def render_chart(path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Patch
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 8, "svg.hashsalt": "task125",
                         "axes.edgecolor": "#8a8984", "axes.linewidth": 0.6})
    fig, ax = plt.subplots(figsize=(7.0, 3.7), dpi=200)
    fig.patch.set_facecolor("#ffffff")
    x = list(range(len(CHART_MONTHS)))
    i0 = CHART_MONTHS.index(PLAN_MONTHS[0])
    ax.axvspan(i0 - 0.5, len(x) - 0.5, color="#ecebe7", zorder=0, lw=0)
    ax.text(i0 + 5.5, 905, "2027/28, the order year (forecast)", ha="center", va="top", fontsize=7.5,
            color="#52514e")
    base = [0] * len(x)
    for dept, col in zip(STACK, COLOURS):
        vals = [HIST[m].get(dept, 0) for m in CHART_MONTHS]
        if not any(vals):
            continue
        alphas = [0.55 if m in PROJECTED else 1.0 for m in CHART_MONTHS]
        for xi, v, b, a in zip(x, vals, base, alphas):
            if v:
                ax.bar(xi, v, bottom=b, width=0.78, color=col, alpha=a, edgecolor="#ffffff", linewidth=0.6,
                       zorder=2)
        base = [b + v for b, v in zip(base, vals)]
    avg = FERN / 12
    ax.axhline(avg, color="#2b2b29", lw=1.1, ls=(0, (4, 2)), zorder=3)
    ax.annotate("Fernhollow basis: %s routed in 2025/26, %s a month" % (fmt(FERN), fmt(int(round(avg)))),
                xy=(4, avg), xytext=(-0.3, 850), fontsize=7.2, color="#2b2b29", va="bottom", ha="left",
                arrowprops=dict(arrowstyle="-", color="#52514e", lw=0.7))
    ib = CHART_MONTHS.index(BUSY)
    vb = A[PLAN_MONTHS.index(BUSY)]
    ax.annotate("%s: %s" % (mlabel(BUSY), fmt(vb)), xy=(ib, vb), xytext=(ib - 4.2, vb + 95), fontsize=7.5,
                color="#0b0b0b", arrowprops=dict(arrowstyle="-", color="#52514e", lw=0.7))
    ax.set_xlim(-0.6, len(x) - 0.4)
    ax.set_ylim(0, 920)
    ticks = [i for i, m in enumerate(CHART_MONTHS) if m[5:] in ("04", "07", "10", "01")]
    ax.set_xticks(ticks)
    ax.set_xticklabels([mlabel(CHART_MONTHS[i]) for i in ticks], fontsize=7)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, p: fmt(int(v))))
    ax.set_ylabel("Payments routed in the month", fontsize=7.5, color="#52514e")
    ax.grid(axis="y", color="#e2e1dc", lw=0.5, zorder=1)
    ax.set_axisbelow(True)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.tick_params(length=2, colors="#52514e")
    handles = [Patch(facecolor=c, edgecolor="none", label=d) for d, c in zip(STACK, COLOURS)]
    ax.legend(handles=handles, loc="upper center", bbox_to_anchor=(0.5, -0.11), ncol=5, frameon=False,
              fontsize=7, handlelength=1.1, columnspacing=1.2)
    ax.set_title("Order for 2027/28: %s routed payments" % fmt(ORDER), loc="left", fontsize=10, fontweight="bold",
                 color="#0b0b0b")
    fig.tight_layout()
    fig.savefig(path, format="png", dpi=200, metadata={"Software": None})
    plt.close(fig)


# ---------------------------------------------------------------------------------------- the order note (docx)
WHEN = dt.datetime(2027, 3, 3, 16, 40)


def _page_field(par):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for kind, text in (("begin", None), (None, " PAGE "), ("end", None)):
        r = par.add_run()
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        r._r.append(el)


def write_docx(path, png):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10)
    st.paragraph_format.space_after = Pt(6)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    h = sec.header.paragraphs[0]
    h.text = "Wealdmoor County Council, Exchequer Services"
    h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    h.runs[0].font.size = Pt(8)
    f = sec.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = f.add_run("WCC/FA/2027-28   Page ")
    r.font.size = Pt(8)
    _page_field(f)

    def para(text, bold=False, size=None, italic=False, after=None):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        return p

    para("2027/28 examination call-off: we order %s routed payments" % fmt(ORDER), bold=True, size=14, after=4)
    meta = [("To", "Abigail McDonald, Director of Finance (s151)"),
            ("Copy", "Shaun Collins, Category Manager, Procurement; Wendy Lyons, Payment Assurance"),
            ("From", "Tina Rogers, Head of Exchequer Services"),
            ("Date", "3 March 2027"),
            ("Call-off", "WCC/FA/2027-28 under PASF-2023-L2 Lot 2, to Fernhollow Assurance Ltd by Friday 12 March 2027")]
    t = doc.add_table(rows=len(meta), cols=2)
    for i, (k, v) in enumerate(meta):
        t.cell(i, 0).text, t.cell(i, 1).text = k, v
        t.cell(i, 0).paragraphs[0].runs[0].bold = True
        for c in (0, 1):
            t.cell(i, c).paragraphs[0].paragraph_format.space_after = Pt(0)
            t.cell(i, c).paragraphs[0].runs[0].font.size = Pt(9)
        t.cell(i, 0).width, t.cell(i, 1).width = Cm(2.2), Cm(14.4)
    para("", after=2)

    para("We will order %s routed payments for examination for 1 April 2027 to 31 March 2028, a quarterly "
         "allocation of %s. That is the filter's forecast routing for the year, %s payments, to the nearest "
         "hundred.¹" % (fmt(ORDER), fmt(ORDER // 4), fmt(ORDER_U)))
    para("Adult Social Care carries the largest share, %s of the %s, all of it in cell 14: direct payments and "
         "Shared Lives carer payments." % (fmt(int(round(TOP_N, -2))), fmt(ORDER)))
    para("Fernhollow will size on %s, the routed count in our 2025/26 run log, which is the indicative volume "
         "clause 4.2 of the call-off terms gives them and what Douglas Harris expects us to order. We sit %s "
         "below it. Our payments have not changed; what the filter looks at in 2027/28, and who is still in "
         "those cells, has. Three things take us from their figure to ours." % (fmt(FERN_H), fmt(GAP)))
    para("The cells. The 2025/26 log ran on the cells notified from the 2023/24 screen, which included Waste & "
         "Environment cell 11 (%s payments routed last year). Under section 5 of the screen methodology the "
         "2027/28 filter runs on the cells of the 2025/26 screen, the latest statement published when we place "
         "the order (July 2026). We rebuilt that screen from the spending file on gross payments of 1,000 to "
         "999,999.99 pounds with credit notes left out; it reproduces all %d figures in Internal Audit's three "
         "statements and flags five cells, none of them Waste & Environment's." % (fmt(WE_2526), REPRODUCED))
    para("Tenancy Sustainment. Housing Support's cell 14 is mostly the 2024-25 round of Tenancy Sustainment "
         "Payments, %d households paid monthly from the month of approval. The round closed to applications on "
         "30 June 2025 and nothing replaces it. Every one of the 288 households in the 2021 round was paid exactly "
         "%d instalments, and on that term the 2024-25 households have %s instalments left in 2027/28 against %s "
         "paid in 2025/26. Pauline Pollard is right that these payments are approved and audited, but the filter "
         "routes every payment in the cell, so the order has to cover them while they last."
         % (LIVE_2425, TERM, fmt(HS_INSTAL), fmt(HS_INST_2526)))
    para("Shared Lives. The Home First step-down places close on 31 March 2027. The step-down payments "
         "themselves never sat in a flagged cell, but a Shared Lives carer is paid one four-weekly amount for "
         "every guest they host, and %d carers host a long-term band 2 guest alongside a step-down guest. They "
         "are paid %s pounds today; from the first 2027/28 run they are paid %s pounds for the band 2 guest "
         "alone, which is cell 14."
         % (SL_MOVED, ", ".join(fmt(a) for a in DUAL_NOW[:-1]) + " or " + fmt(DUAL_NOW[-1]), fmt(DUAL_AFTER[0])))
    para("The same happens one step removed for %d households approved as joint carers. We pay a joint household's fee in two equal halves, one to each carer's vendor number, and the "
         "households paid %s, %s or %s pounds a half host a band 1 guest, a band 3 guest and a step-down guest. "
         "Without the step-down guest each half is %s pounds, also cell 14. Another %d joint households host a "
         "band 2 guest and a step-down guest and are paid %s, %s or %s pounds a half. When a joint household is "
         "left with one guest we stop the second half and pay the whole fee to the first carer, as we did for "
         "the households that lost a guest in %s, %s and %s, so from April each of these is paid %s pounds once "
         "a run, cell 14, not two halves of %s pounds below the filter's 1,000-pound floor. That puts %d "
         "payments in cell 14 on every carer run (%d now) across the %d runs in the 2027/28 payment calendar. "
         "Wendy's point that Shared Lives has paid the same carers the same amounts for three years holds for "
         "every year up to this one."
         % (HALF_HH, fmt(HALF_AMTS[0] // 100), fmt(HALF_AMTS[1] // 100), fmt(HALF_AMTS[2] // 100),
            fmt(HALF_AFTER[0] // 100), WHOLE_HH, fmt(WHOLE_AMTS[0] // 100), fmt(WHOLE_AMTS[1] // 100),
            fmt(WHOLE_AMTS[2] // 100), ONE_GUEST_DATES[0].strftime("%B %Y"), ONE_GUEST_DATES[1].strftime("%B %Y"),
            ONE_GUEST_DATES[2].strftime("%B %Y"), fmt(WHOLE_AFTER[0] // 100), fmt(WHOLE_HALVED[0] // 100),
            SL14_AFTER, SL14_NOW, len(SL_RUNS)))

    rows = [("", "Payments"),
            ("Routed in 2025/26 (run log), Fernhollow's basis", fmt(FERN)),
            ("Waste & Environment cell 11, not flagged on the 2025/26 screen", "-" + fmt(WE_2526)),
            ("Housing Support cell 14, Tenancy Sustainment instalments running off", "-" + fmt(-HS_DELTA)),
            ("Adult Social Care cell 14, Shared Lives carers after the step-down closure", "+" + fmt(SL_DELTA_SINGLE)),
            ("Adult Social Care cell 14, joint carer households' halves after the closure", "+" + fmt(SL_DELTA_JOINT)),
            ("Adult Social Care cell 14, joint households left with one guest, paid whole", "+" + fmt(SL_DELTA_WHOLE)),
            ("Forecast routed in 2027/28", fmt(ORDER_U)),
            ("Order, to the nearest hundred", fmt(ORDER))]
    para("From Fernhollow's figure to the order", bold=True, size=10, after=2)
    tb = doc.add_table(rows=len(rows), cols=2)
    tb.style = "Table Grid"
    tb.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, (a, b) in enumerate(rows):
        tb.cell(i, 0).text, tb.cell(i, 1).text = a, b
        tb.cell(i, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for c in (0, 1):
            pp = tb.cell(i, c).paragraphs[0]
            pp.paragraph_format.space_after = Pt(0)
            for rr in pp.runs:
                rr.font.size = Pt(9)
                rr.bold = i in (0, len(rows) - 1)
        tb.cell(i, 0).width, tb.cell(i, 1).width = Cm(12.6), Cm(2.6)
    para("Source: digit_filter_run_log_2025-26.xlsx; spending file April 2023 to February 2027; Tenancy "
         "Sustainment case extract; Shared Lives rate schedule; 2027/28 BACS calendar.", italic=True, size=8)

    doc.add_picture(io.BytesIO(png), width=Cm(16.4))
    para("Figure 1. Payments routed each month by department, April 2025 to March 2028. 2025/26 is the run log; "
         "April 2026 to January 2027 is counted from the spending file on the cells the 2026/27 call-off used; "
         "February and March 2027 (paler) run past the extract and are projected; 2027/28 is the forecast.",
         italic=True, size=8)

    para("What getting it wrong costs", bold=True, size=10, after=2)
    para("Last year shows both sides. On the original allocation of %s a quarter we sent %s examinations to "
         "the premium rate in the first half of the year; after the variation raised it to %s from October we "
         "paid %s pounds in unused-volume charges in the second half. The quarter-by-quarter figures and the "
         "monthly working are in examination_calloff_2027-28.xlsx."
         % (fmt(ALLOC0), fmt(sum(PREM[:2])), fmt(ALLOC1), fmt(sum(UNUSED_GBP[2:]))))
    para("Tina Rogers", after=0)
    para("Head of Exchequer Services", after=10)
    para("¹ Routed payments are counted in the month the filter runs on them, the month the BACS file is "
         "submitted (run book section 2). Counted by payment date, 1 April 2027 to 31 March 2028, the year comes "
         "to the same %s." % fmt(ORDER_U), size=8)
    cp = doc.core_properties
    cp.author = cp.last_modified_by = "Tina Rogers"
    cp.title = "2027/28 examination call-off"
    cp.comments = cp.subject = cp.keywords = cp.category = ""
    cp.revision = 3
    cp.created = WHEN - dt.timedelta(days=1, hours=5)
    cp.modified = WHEN
    doc.save(str(path))
    W.scrub_ooxml(path, "Tina Rogers", WHEN)


# ---------------------------------------------------------------------------------------- the working (xlsx)
def write_xlsx(path):
    wb = W.xlsx_book(path, "2027/28 examination call-off: working", "Wendy Lyons", "Wealdmoor County Council",
                     WHEN)
    F = dict(title=wb.add_format({"bold": True, "font_size": 12}),
             sub=wb.add_format({"italic": True, "font_color": "#52514e"}),
             hdr=wb.add_format({"bold": True, "bg_color": "#ecebe7", "bottom": 1, "text_wrap": True,
                                "valign": "bottom"}),
             hdr_r=wb.add_format({"bold": True, "bg_color": "#ecebe7", "bottom": 1, "text_wrap": True,
                                  "valign": "bottom", "align": "right"}),
             txt=wb.add_format({}),
             n=wb.add_format({"num_format": "#,##0"}),
             gbp=wb.add_format({"num_format": "£#,##0"}),
             mon=wb.add_format({"num_format": "mmm yyyy", "align": "left"}),
             tot=wb.add_format({"bold": True, "top": 1}),
             tot_n=wb.add_format({"bold": True, "top": 1, "num_format": "#,##0"}),
             tot_gbp=wb.add_format({"bold": True, "top": 1, "num_format": "£#,##0"}),
             note=wb.add_format({"text_wrap": True, "valign": "top"}))

    def mdate(m):
        return dt.datetime(int(m[:4]), int(m[5:]), 1)

    def head(ws, title, sub, cols, widths):
        ws.write(0, 0, title, F["title"])
        ws.write(1, 0, sub, F["sub"])
        for j, (c, w) in enumerate(zip(cols, widths)):
            ws.write(3, j, c, F["hdr"] if j == 0 else F["hdr_r"])
            ws.set_column(j, j, w)
        ws.set_row(3, 30)
        ws.freeze_panes(4, 1)

    def col(j):
        return chr(ord("A") + j)

    # 1. the plan year by run month
    ws = wb.add_worksheet("2027-28 by month")
    depts = ["Adult Social Care", "Housing Support", "Highways & Transport", "Property & Facilities"]
    head(ws, "Payments the filter will route, April 2027 to March 2028",
         "By run month (the month the BACS file is submitted). Cells from the 2025/26 screen.",
         ["Run month"] + depts + ["Routed payments"], [13, 14, 14, 14, 14, 13])
    for i, m in enumerate(PLAN_MONTHS):
        r = 4 + i
        ws.write_datetime(r, 0, mdate(m), F["mon"])
        for j, d in enumerate(depts):
            ws.write_number(r, 1 + j, PLAN_BY[m].get(d, 0), F["n"])
        ws.write_formula(r, 5, "=SUM(B%d:E%d)" % (r + 1, r + 1), F["n"], A[i])
    r = 4 + 12
    ws.write(r, 0, "2027/28", F["tot"])
    for j in range(1, 6):
        v = sum(PLAN_BY[m].get(depts[j - 1], 0) for m in PLAN_MONTHS) if j < 5 else ORDER_U
        ws.write_formula(r, j, "=SUM(%s5:%s16)" % (col(j), col(j)), F["tot_n"], v)
    ws.write(r + 2, 0, "Order (nearest hundred)", F["txt"])
    ws.write_number(r + 2, 5, ORDER, F["n"])
    ws.print_area(0, 0, r + 2, 5)
    ws.repeat_rows(3)

    # 2. 2025/26 routed against examined
    ws = wb.add_worksheet("2025-26 routed v examined")
    head(ws, "2025/26: payments routed and payments Fernhollow examined, by run month",
         "Routed from the run log (a re-run replaces its scheduled run; a supplementary run adds). Examined from "
         "Fernhollow's acknowledgements, against the run each batch came from.",
         ["Run month", "Payments routed", "Payments examined"], [13, 14, 14])
    for i, m in enumerate(M2526):
        r = 4 + i
        ws.write_datetime(r, 0, mdate(m), F["mon"])
        ws.write_number(r, 1, B1[i], F["n"])
        ws.write_number(r, 2, B2[i], F["n"])
    r = 16
    ws.write(r, 0, "2025/26", F["tot"])
    ws.write_formula(r, 1, "=SUM(B5:B16)", F["tot_n"], FERN)
    ws.write_formula(r, 2, "=SUM(C5:C16)", F["tot_n"], sum(B2))
    ws.print_area(0, 0, r, 2)

    # 3. 2025/26 quarters
    ws = wb.add_worksheet("2025-26 by quarter")
    head(ws, "2025/26: examinations at the premium rate and unused-volume charges, by quarter",
         "Call-off WCC/FA/2025-26 as varied from 1 October 2025; 2025/26 rate card; call-off terms clauses 5 and 6.",
         ["Quarter", "Allocation", "Examinations charged", "At the premium rate", "Unused allocation",
          "Unused-volume charge"], [24, 11, 13, 12, 12, 13])
    qn = ["Q1 (Apr to Jun 2025)", "Q2 (Jul to Sep 2025)", "Q3 (Oct to Dec 2025)", "Q4 (Jan to Mar 2026)"]
    for i, q in enumerate((1, 2, 3, 4)):
        r = 4 + i
        ws.write(r, 0, qn[i], F["txt"])
        ws.write_number(r, 1, ALLOC[q], F["n"])
        ws.write_number(r, 2, CHARGED[q], F["n"])
        ws.write_number(r, 3, PREM[i], F["n"])
        ws.write_number(r, 4, UNUSED[i], F["n"])
        ws.write_number(r, 5, UNUSED_GBP[i], F["gbp"])
    r = 8
    ws.write(r, 0, "2025/26", F["tot"])
    for j, v in ((1, sum(ALLOC.values())), (2, sum(CHARGED.values())), (3, sum(PREM)), (4, sum(UNUSED))):
        ws.write_formula(r, j, "=SUM(%s5:%s8)" % (col(j), col(j)), F["tot_n"], v)
    ws.write_formula(r, 5, "=SUM(F5:F8)", F["tot_gbp"], sum(UNUSED_GBP))
    ws.write(r + 2, 0, "Base rate %s per examination; unused allocation charged at %d%% of it (clause 6.3). "
             "Batches under %d examinations charged as %d (clause 5.2); withdrawn batches and returned payments "
             "not charged (5.3)." % ("£%.2f" % BASE, int(SHARE * 100), MIN_BATCH, MIN_BATCH), F["sub"])
    ws.print_area(0, 0, r + 2, 5)

    # notes
    ws = wb.add_worksheet("Notes")
    ws.set_column(0, 0, 22)
    ws.set_column(1, 1, 90)
    notes = [("Prepared", "Wendy Lyons, Payment Assurance, for Tina Rogers. As at 3 March 2027."),
             ("Spending file", "wealdmoor_spend_over_500_2023-04_to_2027-02.csv, payments to the 25 February "
                               "2027 run."),
             ("Screen", "Gross payments (net plus VAT) of 1,000.00 to 999,999.99 pounds, credit notes left out, "
                        "first two digits by department and year; reproduces all 63 published statement figures."),
             ("Cells", "2027/28 routes on the 2025/26 screen's cells (methodology section 5): Adult Social Care 14, "
                       "Housing Support 14, Highways & Transport 49 and 99, Property & Facilities 12."),
             ("Run month", "The month the BACS file is submitted (run book section 2; dates from "
                           "bacs_payment_calendar_2025-26_to_2027-28.xlsx)."),
             ("Tenancy Sustainment", "Each household paid %d monthly instalments from the month of approval, as "
                                     "every 2021 household was." % TERM),
             ("Shared Lives", "Each carer's four-weekly amount re-summed over the placements left after the Home "
                              "First step-down closure on 31 March 2027, at the scheduled weekly rates. A jointly "
                              "approved household is paid in two equal halves, one to each carer's vendor number, "
                              "while it hosts two or more guests, and its whole fee to the first carer while it "
                              "hosts one, as the households that lost a guest in the extract were paid.")]
    for i, (k, v) in enumerate(notes):
        ws.write(i, 0, k, F["hdr"])
        ws.write(i, 1, v, F["note"])
    wb.worksheets()[0].activate()
    W.close_xlsx(wb, path, "Wendy Lyons", WHEN)


# ---------------------------------------------------------------------------------------- run
if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    buf = io.BytesIO()
    render_chart(buf)
    write_docx(OUT / "examination_calloff_2027-28.docx", buf.getvalue())
    write_xlsx(OUT / "examination_calloff_2027-28.xlsx")
    for p in OUT.iterdir():
        W.set_mtime(p, WHEN)
    print("Screen: gross, complete decades, credit notes excluded, reproduces %d of 63 statement figures" % REPRODUCED)
    print("2025/26 cells applied in 2027/28: %s" % ", ".join("%s %d" % (SHORT[d], c) for d, c in sorted(PLAN_CELLS)))
    print("Housing Support: term %d instalments (2021 round %s); 2027/28 instalments %s; cell 14 %s"
          % (TERM, dict(TERMS_SEEN), fmt(HS_INSTAL), fmt(BY_CELL[("Housing Support", 14)])))
    print("Shared Lives: %d carer payments a run in cell 14 now, %d after the closure (%d single carers and %d "
          "halves of %d joint households and %d one-guest households paid whole moving in), %d runs; ASC cell 14 %s"
          % (SL14_NOW, SL14_AFTER, SL_MOVED, HALF_MOVED, HALF_HH, WHOLE_MOVED, len(SL_RUNS),
             fmt(BY_CELL[("Adult Social Care", 14)])))
    print("Joint households %d; halves moving in %s, after the closure %s; one-guest households paid whole %d "
          "(halves %s, whole %s); one-guest changes in the extract %s" % (JOINT_HH, HALF_AMTS, HALF_AFTER, WHOLE_HH,
                                                                          WHOLE_AMTS, WHOLE_AFTER, ONE_GUEST_DATES))
    print("Routed 2027/28: %s, order %s; largest share %s %s (%s)" % (fmt(ORDER_U), fmt(ORDER), TOP_DEPT,
                                                                      fmt(TOP_N), fmt(int(round(TOP_N, -2)))))
    print("Fernhollow's figure %s (%s), gap %s; rung 3 (carers at current amounts) %s; rung 4 (halves held) %s; "
          "rung 5 (every joint household halved) %s" % (fmt(FERN), fmt(FERN_H), fmt(GAP), fmt(RUNG3), fmt(RUNG4),
                                                         fmt(RUNG5)))
    print("A " + json.dumps(A))
    print("B1 " + json.dumps(B1))
    print("B2 " + json.dumps(B2))
    print("C premium " + json.dumps(PREM) + " unused GBP " + json.dumps(UNUSED_GBP))
