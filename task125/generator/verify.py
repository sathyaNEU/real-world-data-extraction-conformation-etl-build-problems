"""Independent verifier: recomputes the July to September split, the gap path and the three regional asks from target/
alone, with its own parsing, raking and matching (no generator module), and compares them with the goldens."""
import csv
import os
import re
import sys
import zipfile
from collections import defaultdict

import numpy as np
import openpyxl
import pyarrow.parquet as pq

REG = ["VEL", "CAR", "PEL", "HAS", "TIR", "DUN"]
NAME = {"VEL": "Velmouth", "CAR": "Carrowdene", "PEL": "Pellard", "HAS": "Haskvale", "TIR": "Tirrin", "DUN": "Dunmarrow"}


def controls(target):
    wb = openpyxl.load_workbook(os.path.join(target, "population_controls_2023q1_2027q4.xlsx"), data_only=True)
    reg = {r[0]: np.array(r[2:8], float) for r in wb["Region"].iter_rows(min_row=2, values_only=True)}
    cel = {r[0]: np.array(r[2:16], float) for r in wb["Age and sex"].iter_rows(min_row=2, values_only=True)}
    return reg, cel


def ipf(region, cell, rt, ct):
    w = np.ones(len(region))
    for _ in range(500):
        w = w * rt[region] / np.bincount(region, w, 6)[region]
        f = ct / np.bincount(cell, w, 14)
        w = w * f[cell]
        if abs(f - 1).max() < 1e-13:
            break
    return w


def ur_er(st, w):
    e, u = w[st == 1].sum(), w[st == 2].sum()
    return 100 * u / (e + u), 100 * e / w.sum()


def main(target, golden):
    t = pq.read_table(os.path.join(target, "hls_person_quarter_2023q1_2026q3.parquet")).to_pydict()
    n = len(t["quarter"])
    first = {}
    with open(os.path.join(target, "hls_call_records_2024q1_2026q3.csv"), encoding="utf-8") as fh:
        calls = list(csv.DictReader(fh))
    for r in calls:
        if r["outcome"] == "110":
            s = r["serial"][:9]
            first[s] = min(first.get(s, 99), int(r["visit_no"]))
    reg, cel = controls(target)

    def sample(q):
        idx = [i for i in range(n) if t["quarter"][i] == q]
        region = np.array([REG.index(t["region"][i]) for i in idx])
        age = np.array([t["age"][i] for i in idx])
        grp = np.digitize(age, [25, 35, 45, 55, 65, 75])
        cell = 2 * grp + np.array([t["sex"][i] for i in idx]) - 1
        st = np.array([t["ilo_status"][i] for i in idx])
        wave = np.array([t["wave"][i] for i in idx])
        late = np.array([first.get(t["serial"][i][:9], 0) >= 4 for i in idx])
        wt = np.array([t["weight"][i] for i in idx])
        return region, cell, st, wave, late, wt

    region, cell, st, wave, late, _ = sample("2026Q2")
    base = ur_er(st, ipf(region, cell, reg["2026Q2"], cel["2026Q2"]))
    path = []
    for pos in [(1,), (2, 1), (3, 2), (4, 3), (5, 4)]:
        keep = ~(np.isin(wave, pos) & late)
        r = ur_er(st[keep], ipf(region[keep], cell[keep], reg["2026Q2"], cel["2026Q2"]))
        path.append((r[0] - base[0], r[1] - base[1]))
    q2 = ur_er(*sample("2026Q2")[2::3])
    q3 = ur_er(*sample("2026Q3")[2::3])
    real_u = q3[0] - path[0][0] - q2[0]

    # asks
    register = {}
    with open(os.path.join(target, "hls_address_register_2024q1_2026q3.csv"), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            register[r["serial"]] = r["region"]
    vis = defaultdict(set)
    for r in calls:
        if r["intake"] == "2026Q3":
            vis[register[r["serial"][:9]]].add((r["serial"], r["visit_no"]))
    key = lambda i: (t["serial"][i][:9], t["sex"][i], str(t["date_of_birth"][i]))
    q2p = {key(i): i for i in range(n) if t["quarter"][i] == "2026Q2"}
    q3p = {key(i): i for i in range(n) if t["quarter"][i] == "2026Q3"}
    num, den = defaultdict(int), defaultdict(int)
    for k, i in q2p.items():
        if t["ilo_status"][i] == 2 and k in q3p:
            den[t["region"][i]] += 1
            num[t["region"][i]] += t["ilo_status"][q3p[k]] == 1
    hours = defaultdict(float)
    for i in range(n):
        if t["quarter"][i] == "2026Q3" and t["ilo_status"][i] == 1:
            h = t["actual_hours"][i]
            if t["mode"][i] == "CATI":
                whole = int(h)
                h = whole + round((h - whole) * 100) / 60
            hours[t["region"][i]] += t["weight"][i] * h

    # compare with the goldens
    wb = openpyxl.load_workbook(os.path.join(golden, "hls_visit_limit_path_2027.xlsx"), data_only=True)
    gp = [r for r in wb["Gap path"].iter_rows(min_row=5, max_row=9, values_only=True)]
    for (u, e), r in zip(path, gp):
        assert "%.2f" % r[1] == "%.2f" % round(u, 2) and "%.2f" % r[2] == "%.2f" % round(e, 2), (r, u, e)
    split = {r[0]: r for r in wb["Gap path"].iter_rows(min_row=12, max_row=17, values_only=True)}
    assert "%.2f" % split["of which the labour market"][1] == "%.2f" % round(real_u, 2)
    assert "%.2f" % split["Published, July to September 2026"][1] == "%.2f" % round(q3[0], 2)
    rg = {r[0]: r for r in wb["Regions"].iter_rows(min_row=4, max_row=9, values_only=True)}
    for g in REG:
        row = rg[NAME[g]]
        assert row[1] == len(vis[g]), (g, row[1], len(vis[g]))
        assert (row[2], row[3]) == (den[g], num[g]), (g, row[2:4], den[g], num[g])
        assert "%.2f" % row[4] == "%.2f" % (100 * num[g] / den[g])
        assert "%.1f" % row[5] == "%.1f" % (hours[g] / 1000), (g, row[5], hours[g])
    with zipfile.ZipFile(os.path.join(golden, "hls_visit_limit_note_2026q3.docx")) as z:
        x = re.sub(r"<[^>]+>", "", z.read("word/document.xml").decode("utf-8"))
    assert "accounts for %.2f point of that rise and the labour market for %.2f point" % (round(path[0][0], 2),
                                                                                          round(real_u, 2)) in x
    print("verify: OK (July to September: %.2f of the %.2f-point rise is the cap; gap path %s; 18 regional figures "
          "recomputed independently)" % (path[0][0], q3[0] - q2[0], "/".join("%.2f" % p[0] for p in path)))


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    task = os.path.dirname(here)
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(task, "target"),
         sys.argv[2] if len(sys.argv) > 2 else os.path.join(task, "golden"))
