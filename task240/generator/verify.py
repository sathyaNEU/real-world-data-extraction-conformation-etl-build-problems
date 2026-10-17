"""task240 independent verifier: recomputes every owed figure from the shipped bundle alone (its own code, nothing from
the generator), then checks figures.json, the golden workbook and note, and the submission against it."""
import json
import os
import sys

import numpy as np
import pandas as pd

LINE = 20.0
IND = ["PEN", "RENT", "CARE", "ARR"]
NAMES = {"PEN": "Adults receiving a public pension", "RENT": "Adults in privately rented homes",
         "CARE": "Adults providing unpaid care", "ARR": "Adults in households behind on a household bill"}


def _f(target, pat):
    hits = [f for f in os.listdir(target) if pat in f]
    assert len(hits) == 1, (pat, hits)
    return os.path.join(target, hits[0])


def adults(target):
    hh = pd.read_csv(_f(target, "households.csv"))
    mem = pd.read_csv(_f(target, "household_members.csv"))
    wt = pd.read_csv(_f(target, "final_weights.csv"))
    a = mem[mem["age"] >= 16].merge(wt, on=["hh_id", "person_no"]).merge(hh, on="hh_id")
    assert len(a) == len(wt)
    a["y_PEN"] = (a["pension"] == 1).astype(float)
    a["y_RENT"] = (a["tenure"] == 3).astype(float)
    a["y_CARE"] = (a["care"] == 1).astype(float)
    a["y_ARR"] = (a["arrears"] == 1).astype(float)
    a["agegrp"] = pd.cut(a["age"], [15, 29, 49, 64, 200], labels=False)
    a["cell"] = a["region"] + "/" + a["agegrp"].astype(str) + "/" + a["sex"].astype(str)
    return a, hh


def cvs(a):
    """the post-stratified total's linearisation: w (y - cell rate), PSU totals, with-replacement within strata"""
    out = {}
    for y in IND:
        v = a["y_" + y].values
        w = a["final_weight"].values
        rate = (pd.Series(w * v).groupby(a["cell"].values).sum() / pd.Series(w).groupby(a["cell"].values).sum())
        z = w * (v - a["cell"].map(rate).values)
        t = pd.DataFrame({"z": z, "h": a["stratum"].values, "p": a["psu"].values}).groupby(["h", "p"])["z"].sum()
        t = t.reset_index()
        t["reg"] = t["h"].str[:3]
        var = {}
        for (reg, h), g in t.groupby(["reg", "h"]):
            n = len(g)
            var[reg] = var.get(reg, 0.0) + n / (n - 1) * ((g["z"] - g["z"].mean()) ** 2).sum()
        tot = pd.Series(w * v).groupby(a["region"].values).sum()
        for reg in tot.index:
            out[(reg, y)] = (100 * np.sqrt(var[reg]) / tot[reg], tot[reg])
    return out


def asks(target, hh):
    iss = pd.read_csv(_f(target, "issued_sample.csv"))
    out = pd.read_csv(_f(target, "outcomes_"))
    cr = pd.read_csv(_f(target, "call_records.csv"), dtype={"result": str})
    region = dict(zip(iss["psu"], iss["region_code"]))
    region.update(dict(zip(hh["psu"], hh["region"])))           # the PSU's region as the published tables have it
    last = out.sort_values(["address_id", "issue_no"]).drop_duplicates("address_id", keep="last")
    fin = dict(zip(last["address_id"], last["outcome_code"]))
    for h in hh["hh_id"]:
        fin[h] = 110                                             # interviewed, whatever the extract said
    rows = pd.DataFrame({"address_id": iss["address_id"], "region": iss["psu"].map(region),
                         "o": [fin.get(x, 0) for x in iss["address_id"]]})
    elig = rows["o"].isin([110, 210, 220, 230])
    inel = rows["o"].isin([410, 420, 430])
    e = elig.sum() / (elig.sum() + inel.sum())
    rr = {}
    for reg, g in rows.groupby("region"):
        el = g["o"].isin([110, 210, 220, 230])
        un = ~el & ~g["o"].isin([410, 420, 430])
        rr[reg] = 100 * (g["o"] == 110).sum() / (el.sum() + e * un.sum())
    c = cr[cr["event_type"].isin(["VISIT", "PHONE"])]
    c = c.drop_duplicates(["address_id", "issue_no", "event_time", "event_type", "result", "interviewer_id"])
    c = c[c["address_id"].isin(set(hh["hh_id"]))]
    reg_of = dict(zip(hh["hh_id"], hh["region"]))
    calls = c["address_id"].map(reg_of).value_counts().to_dict()
    return rr, {k: int(v) for k, v in calls.items()}


def main(target, golden, figures, submission):
    F = json.load(open(figures))
    sub = open(submission).read()
    a, hh = adults(target)
    C = cvs(a)
    warned = sorted(k for k, (cv, _) in C.items() if cv >= LINE)
    assert sorted("%s %s" % k for k in warned) == sorted(F["call"]["warned"]), warned
    for (reg, y), (cv, tot) in C.items():
        assert abs(round(cv, 1) - F["deliverables"]["cv"][reg][y]) < 1e-9, (reg, y, cv)
        assert abs(round(tot / 1000.0, 1) - F["deliverables"]["estimates_thousands"][reg][y]) < 1e-9, (reg, y)
        x = cv * 10
        assert 0.5 - abs(x - round(x)) >= 0.07, ("CV near its rounding edge", reg, y, cv)
    rr, calls = asks(target, hh)
    for reg in rr:
        assert abs(round(rr[reg], 1) - F["deliverables"]["response_rate"][reg]) < 1e-9, (reg, rr[reg])
        assert calls[reg] == F["deliverables"]["calls"][reg], (reg, calls[reg])
    assert F["steps"]["adults"] == len(a) and F["steps"]["households"] == len(hh)
    # the goldens carry the same figures
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(golden, "hrs27_quality_report.xlsx"), data_only=True)
    ws = wb["Estimates"]
    n = 0
    for r in range(5, 53):
        reg, label = ws.cell(row=r, column=1).value, ws.cell(row=r, column=3).value
        y = [k for k, v in NAMES.items() if v == label][0]
        cv = round(C[(reg, y)][0], 1)
        assert abs(ws.cell(row=r, column=5).value - cv) < 1e-9, (reg, y)
        assert (ws.cell(row=r, column=6).value == "E") == ((reg, y) in warned), (reg, y)
        n += 1
    assert n == 48
    ws = wb["Fieldwork"]
    for r in range(5, 17):
        reg = ws.cell(row=r, column=1).value
        assert abs(ws.cell(row=r, column=5).value - round(rr[reg], 1) / 100.0) < 1e-12, reg
        assert ws.cell(row=r, column=6).value == calls[reg], reg
    assert ws.cell(row=17, column=6).value == sum(calls.values())
    assert ws.cell(row=17, column=3).value == F["steps"]["issued"] and ws.cell(row=17, column=4).value == len(hh)
    import docx
    d = docx.Document(os.path.join(golden, "hrs27_quality_note.docx"))
    cells = [c.text for t in d.tables for row in t.rows for c in row.cells]
    for reg, y in warned:
        assert any(reg in c for c in cells) and "%.1f" % C[(reg, y)][0] in cells, (reg, y)
    # every owed figure appears in the submission's answer block
    for reg, y in C:
        assert "%.1f" % round(C[(reg, y)][0], 1) in sub
    for reg in rr:
        assert "%.1f, %s calls" % (round(rr[reg], 1), "{:,}".format(calls[reg])) in sub, reg
    print("verify: every figure recomputes from the bundle (%d adults, %d households; %d warned; 12 response rates and "
          "12 call counts)" % (len(a), len(hh), len(warned)))


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    task = os.path.dirname(here)
    main(os.path.join(task, "target"), os.path.join(task, "golden"), os.path.join(here, "figures.json"),
         os.path.join(task, "submission.md"))
    sys.exit(0)
