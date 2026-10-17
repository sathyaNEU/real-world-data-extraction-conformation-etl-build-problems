"""task240 documents: the Branch's statistical quality standard (pdf), the agency's weighting note and fieldwork
close-out (docx), the questionnaire codebook (xlsx), the data dictionary and the readme (md)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess

BUREAU = P.CLIENT
AGENCY = P.AGENCY
PEOPLE = {"lead": "Kristina Stephens", "bureau": "Latasha Evans", "field": "Anna Mckee"}
OUTCOMES = [
    (110, "Full interview: every adult in the household interviewed, in person or by proxy", "eligible"),
    (210, "Refusal", "eligible"),
    (220, "Non-contact, the address known to be occupied", "eligible"),
    (230, "Other non-response at an occupied address", "eligible"),
    (310, "No contact; whether a household lives at the address not established", "not established"),
    (410, "Vacant", "not eligible"),
    (420, "Not a private residence", "not eligible"),
    (430, "Demolished or not found", "not eligible"),
]


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.4, leading=13, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=14, leading=18, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=14, spaceBefore=8,
                        spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10.5, textColor=colors.HexColor("#444444"))
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=8.2, leading=10)
    return dict(body=body, h1=h1, h2=h2, small=small, cell=cell)


def _footer(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(18 * mm, 10 * mm, text)
        canvas.drawRightString(letter[0] - 18 * mm, 10 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _table(rows, widths, size=8.2):
    t = Table(rows, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", size), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", size),
                           ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA0AB")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E2E7F0")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


def _p(text, style):
    return Paragraph(text, style)


# ------------------------------------------------------------------ the quality standard (pdf)
def standard(path):
    S = _styles()
    b, h1, h2, sm = S["body"], S["h1"], S["h2"], S["small"]
    st = [_p("Statistical quality standard QS-7: contracted household surveys", h1),
          _p("%s, Ministry of Finance. Edition 3, issued 14 March 2026; applies to every household survey wave the Branch "
             "commissions under contract from the date of issue. Owner: %s, Survey Statistics Lead." % (BUREAU, PEOPLE["bureau"]),
             sm),
          Spacer(1, 4),
          _p("1. Purpose and scope", h2),
          _p("This standard sets how the published estimates of a household survey the Branch commissions are judged "
             "for sampling error, how the survey's response is measured, and what the contractor's quality report for "
             "each wave sets out.", b),
          _p("2. Published estimates", h2),
          _p("2.1 A table estimate is the estimated number of adults (aged 16 and over) in a health region who have the "
             "characteristic the table reports, made with the wave's final weights as the contractor's weighting note "
             "for the wave describes. Estimates are published in thousands.", b),
          _p("2.2 Estimates are published for the twelve health regions as drawn by the Branch's 2025 boundary "
             "review.", b),
          _p("2.3 An adult whose answer was don't know, or who refused the question, is not counted as having the "
             "characteristic. An answer a household member gave for an adult who was not available (a proxy answer) is "
             "that adult's answer.", b),
          _p("3. Sampling error", h2),
          _p("3.1 The coefficient of variation (CV) of a published estimate is its standard error as a percentage of the "
             "estimate, both taken before the estimate is rounded for publication.", b),
          _p("3.2 The standard error is that of the estimate as published (2.1). It is estimated by Taylor "
             "linearisation over the wave's strata and primary sampling units (PSUs) as the contractor's analysis files "
             "give them, the PSUs treated as selected with replacement within their strata.", b),
          _p("3.3 A published estimate whose CV is 20.0 per cent or more carries the low-reliability warning: the letter "
             "E beside the figure in the table.", b),
          _p("4. Response", h2),
          _p("4.1 Each address issued for interview has one final outcome, the outcome of its last issue: an address "
             "reissued to a second interviewer takes the outcome of the reissue. Outcomes are coded as in Table 1.", b)]
    rows = [["Code", "Outcome", "Eligibility"]]
    for code, text, el in OUTCOMES:
        rows.append([str(code), Paragraph(text, S["cell"]), el])
    st += [_table(rows, [14 * mm, 122 * mm, 30 * mm]), _p("Table 1. Final outcome codes.", sm),
           _p("4.2 A health region's response rate is its addresses with a full interview as a percentage of its eligible "
              "addresses: those found eligible, together with the share e of its addresses whose eligibility was not "
              "established. e is the proportion found eligible among all of the wave's addresses whose eligibility was "
              "established, across the twelve regions. Every address issued counts in the rate of the health region it "
              "lies in.", b),
           _p("4.3 A call is a visit to an address, or a telephone call to its household, by the interviewer working the "
              "issue. An advance letter, and a calling card left at the door, are not calls. The calls to an address are "
              "those of every issue.", b),
           _p("5. The quality report and the quality note", h2),
           _p("5.1 With each wave's tables the contractor delivers a quality report setting out every published estimate "
              "with its CV and whether it carries the low-reliability warning, and for each health region the response "
              "rate and the number of calls made to its addresses with a full interview.", b),
           _p("5.2 The Branch publishes a quality note with the tables, naming every estimate that carries the "
              "warning.", b),
           _p("6. Review", h2),
           _p("The Branch reviews this standard every two years. The next review is due in March 2028.", b)]
    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title="QS-7 Statistical quality standard: contracted household surveys",
                            author=PEOPLE["bureau"], creator=BUREAU, invariant=1)
    foot = "%s | QS-7 Statistical quality standard | Edition 3 (March 2026)" % BUREAU
    doc.build(st, onFirstPage=_footer(foot), onLaterPages=_footer(foot))
    postprocess.scrub(path, BUREAU, "2026-03-14")


# ------------------------------------------------------------------ docx helpers
def _docx(title_line, meta, paras, path, author, created, modified, title, sign=None):
    import docx
    from docx.shared import Pt
    doc = docx.Document()
    stl = doc.styles["Normal"]
    stl.font.name = "Calibri"
    stl.font.size = Pt(10.5)
    head = doc.add_paragraph()
    r = head.add_run(title_line)
    r.bold = True
    r.font.size = Pt(12)
    for k, v in meta:
        p = doc.add_paragraph()
        a = p.add_run(k + ": ")
        a.bold = True
        p.add_run(v)
        p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph()
    for t in paras:
        if isinstance(t, tuple):
            p = doc.add_paragraph()
            a = p.add_run(t[0])
            a.bold = True
            p.add_run(" " + t[1])
        else:
            p = doc.add_paragraph(t)
        p.paragraph_format.space_after = Pt(6)
    if sign:
        p = doc.add_paragraph(sign)
        p.paragraph_format.space_before = Pt(8)
    doc.core_properties.author = author
    doc.core_properties.title = title
    doc.save(path)
    postprocess.scrub(path, author, created[:10])
    postprocess.finish_office(path, created, modified, AGENCY)


def weighting_note(path, F, f):
    paras = [
        "The final weights for the 2027 wave were frozen on 12 November 2027 and are in %s. They were made in two "
        "steps, as for the 2025 and 2026 waves." % F["weights"],
        ("Design weights.", "The wave's %s addresses came from the Branch's master sample: PSUs selected within strata "
                            "formed by health region and by urban and rural areas, then about a dozen addresses selected "
                            "in each selected PSU. Each address's design weight is the inverse of its probability of "
                            "selection, and every adult in the household at the address takes it." % f["issued"]),
        ("Post-stratification.", "The design weights of the adults in households giving a full interview were then "
                                 "scaled, within each health region, separately for each age group (16 to 29, 30 to 49, "
                                 "50 to 64, 65 and over) and sex, so that in each of these 96 cells they add to the "
                                 "Branch's estimate of the cell's adult population on 1 July 2027. The Branch publishes "
                                 "its estimates by five-year age band, and the bands were added together to give the four "
                                 "groups. An adult's age group is taken from their age at interview. Every cell held "
                                 "responding adults, so no cells were combined, and no weight was trimmed."),
        ("Final weights.", "The final weight is the post-stratified weight, so the final weights add to the Branch's adult "
                           "population in every region, age group and sex. Only adults (16 and over) are weighted; "
                           "children on a household's roster carry no weight."),
        ("Files.", "%s gives each responding household's health region, stratum and PSU and its household answers; %s "
                   "each member's age, sex and own answers; %s each adult's final weight, keyed by household and person "
                   "number." % (F["households"], F["members"], F["weights"])),
    ]
    _docx("%s | %s 2027" % (AGENCY, P.SURVEY),
          [("Title", "Weighting the 2027 wave"), ("Author", "%s, Lead Survey Statistician" % PEOPLE["lead"]),
           ("To", "%s, Survey Statistics Lead, %s" % (PEOPLE["bureau"], BUREAU)), ("Date", "12 November 2027")],
          paras, path, PEOPLE["lead"], "2027-11-11T19:40:00Z", "2027-11-12T15:05:00Z", "Weighting the 2027 wave")


def close_out(path, F, f):
    paras = [
        "Fieldwork for the 2027 wave ran from 7 April to 26 September. %s interviewers worked the wave's %s PSUs, each "
        "PSU's addresses assigned to one interviewer, with assignment start dates spread from the opening to "
        "mid-August." % (f["interviewers_word"], f["psus"]),
        ("Advance letters.", "The office posted an advance letter to each address about a week before its interviewer's "
                             "assignment started, and again before a reissue."),
        ("Calls.", "Interviewers call in person and, where a household has given a number, by telephone, at different "
                   "times of day and on different days, and leave a calling card when nobody answers the door."),
        ("Reissues.", "From mid-May we reissued addresses where the first interviewer got no interview to a second "
                      "interviewer in the same region, as interviewer time allowed; the last reissues went out in "
                      "mid-September."),
        ("Outcomes.", "The outcome file with this note (%s) was extracted from the case system on 30 September." % F["outcomes"]),
        ("Next.", "Data checks and coding run through October; the weights follow once the Branch's population "
                  "estimates for 1 July are in."),
    ]
    _docx("%s | %s 2027" % (AGENCY, P.SURVEY),
          [("To", "%s, Lead Survey Statistician" % PEOPLE["lead"]), ("From", "%s, Fieldwork Manager" % PEOPLE["field"]),
           ("Date", "4 October 2027"), ("Subject", "Household Resources Survey 2027: fieldwork close-out")],
          paras, path, PEOPLE["field"], "2027-10-04T14:10:00Z", "2027-10-04T20:35:00Z",
          "Household Resources Survey 2027: fieldwork close-out", sign=PEOPLE["field"])


# ------------------------------------------------------------------ the codebook (xlsx)
QUESTIONS = [
    ("tenure", "households", "Does your household own or rent this accommodation?", "the household (one answer)",
     "1 owned outright; 2 owned with a mortgage or loan; 3 rented from a private landlord; 4 rented from a social or "
     "non-profit housing provider; 5 rent-free or other; -8 don't know"),
    ("arrears", "households", "In the last twelve months, has your household been behind with paying any household bill "
                              "(rent or mortgage, electricity, heating, water, phone or internet)?",
     "the household (one answer)", "1 yes; 2 no; -9 refused"),
    ("pension", "members", "Do you receive a public pension (Old Age Security, the Canada Pension Plan or a provincial "
                           "pension)?", "adults aged 55 and over",
     "1 yes; 2 no; -1 not asked; -8 don't know; -9 refused"),
    ("care", "members", "Do you look after, or give any help or support to, anyone because of a long-term physical or "
                        "mental health condition, illness or old age, other than as part of paid work?",
     "adults aged 16 and over", "1 yes; 2 no; -1 not asked; -8 don't know; -9 refused"),
]
TABLES = [
    ("PEN", "Adults receiving a public pension", "members: pension = 1"),
    ("RENT", "Adults in privately rented homes", "households: tenure = 3, each adult in the household"),
    ("CARE", "Adults providing unpaid care", "members: care = 1"),
    ("ARR", "Adults in households behind on a household bill", "households: arrears = 1, each adult in the household"),
]


def codebook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    ws.append(["variable", "file", "question", "asked of", "codes"])
    for row in QUESTIONS:
        ws.append(list(row))
    w2 = wb.create_sheet("Published tables")
    w2.append(["table", "estimate", "counted from"])
    for row in TABLES:
        w2.append(list(row))
    w3 = wb.create_sheet("Call log codes")
    w3.append(["column", "code", "meaning"])
    for code, text in (("VISIT", "the interviewer called at the address"), ("PHONE", "the interviewer telephoned the household"),
                       ("LETTER", "an advance letter posted by the office"),
                       ("CARD", "a calling card left at the door")):
        w3.append(["event_type", code, text])
    for code, text in (("IC", "interview complete"), ("AP", "appointment made"), ("RF", "refusal"),
                       ("NC", "no contact"), ("OT", "other outcome at an occupied address"), ("VA", "vacant"),
                       ("NR", "not a private residence"), ("DM", "demolished or not found")):
        w3.append(["result", code, text])
    w4 = wb.create_sheet("Outcome codes")
    w4.append(["outcome_code", "meaning"])
    w4.append(["110 to 430", "final outcome codes as in QS-7 Table 1"])
    for sh, widths in ((ws, [11, 12, 70, 24, 60]), (w2, [8, 44, 50]), (w3, [12, 9, 44]), (w4, [14, 50])):
        for c in sh[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E2E7F0")
        for i, w in enumerate(widths, 1):
            sh.column_dimensions[get_column_letter(i)].width = w
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        sh.freeze_panes = "A2"
    wb.save(path)
    postprocess.scrub(path, PEOPLE["lead"], "2027-03-19")
    postprocess.finish_office(path, "2027-03-17T16:20:00Z", "2027-03-19T13:45:00Z", AGENCY)


# ------------------------------------------------------------------ the dictionary and the readme
def dictionary(F):
    return """# Data dictionary

{survey} 2027 files held by {agency}. Times are the province's local time.

## {households}
One row per household that gave a full interview (final outcome 110).
- `hh_id`: the household's identifier, the address identifier of its address.
- `region`: the household's health region.
- `stratum`: the stratum of the household's PSU for estimation: the health region, then `U` (urban) or `R` (rural).
- `psu`: the primary sampling unit.
- `interview_date`: the date of the interview.
- `adults`, `children`: the household's members aged 16 and over, and under 16.
- `tenure`, `arrears`: the household's answers (codebook, sheet Questions).

## {members}
One row per member on each responding household's roster: `hh_id`; `person_no` (the member's number on the household's
roster, from 1); `age`; `sex` (1 male, 2 female); `proxy` (`Y` where another member answered for the adult; blank
for children); `pension` and `care` (codebook, sheet Questions).

## {weights}
One row per adult in a responding household: `hh_id`, `person_no`, `final_weight` (the weighting note).

## {issued}
One row per address issued for interview in 2027.
- `address_id`: the address identifier.
- `psu`: the address's PSU.
- `region_code`: the health region code held for the PSU in the master sample.
- `psu_frame`: the master-sample selection the PSU came from: `2023` (the 2023 selection) or `2026` (the 2026 top-up).
- `dwelling`: the dwelling type recorded at listing.

## {outcomes}
The case system's outcomes as extracted on 30 September 2027: one row per issue of an address. `address_id`;
`issue_no` (1 for the first issue, 2 for a reissue); `interviewer_id`; `outcome_code` (QS-7 Table 1);
`outcome_date` (the date of the call the outcome comes from).

## {calls}
The call records exported on 18 October 2027: one row per record received from interviewers' tablets and the
office.
- `record_id`: the office's number for the record, given as each record arrived.
- `address_id`, `issue_no`.
- `event_no`: the record's number among its address's records for the issue, as the tablet wrote them.
- `event_time`: when the call was made, the letter posted or the card left.
- `event_type`, `result`: codebook, sheet Call log codes.
- `interviewer_id`: the interviewer working the issue (`OFFICE` for letters).
- `received_at`: when the record reached the office.

## {interviewers}
The wave's interviewer list: `interviewer_id`, `home_region`, `psus_assigned`, `hours_claimed` (hours claimed on
timesheets for the wave).

## {population}
The Branch's population estimates of adults on 1 July 2027 by health region, sex and five-year age band, in thousands.

## {codebook}
The questions behind the published tables and their codes, the published tables, and the call log codes.
""".format(survey=P.SURVEY, agency=AGENCY, **F)


def readme(F):
    return """# {survey} 2027: wave files

{agency}'s files for the 2027 wave of the {survey}, run for the {bureau} under contract, gathered by
{lead} (Lead Survey Statistician) on Friday 3 December 2027 for the wave's quality report. The tables and the quality
report are due to the Branch on 10 December.

- `{standard}`: the Branch's statistical quality standard QS-7, edition 3.
- `{note}`: how the 2027 weights were made.
- `{close}`: the fieldwork close-out, 4 October 2027.
- `{codebook}`: the 2027 codebook.
- `{households}`, `{members}`, `{weights}`: the analysis files, frozen with the weights on 12 November 2027.
- `{issued}`: the issued sample.
- `{outcomes}`: the outcome extract of 30 September 2027.
- `{calls}`: the call records exported on 18 October 2027.
- `{interviewers}`: the interviewer list.
- `{population}`: the Branch's population estimates for 1 July 2027.
- `DATA_DICTIONARY.md`: the files and their fields.
""".format(survey=P.SURVEY, agency=AGENCY, bureau=BUREAU, lead=PEOPLE["lead"], **F)
