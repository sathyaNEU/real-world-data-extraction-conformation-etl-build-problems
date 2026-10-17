"""task240 golden deliverables: the 2027 quality report workbook and the quality note the Branch publishes."""
import os

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import params as P
import postprocess
import writeup as WU

XLSX = "hrs27_quality_report.xlsx"
DOCX = "hrs27_quality_note.docx"
NAVY, RUST, INK = "23395B", "B5482B", "1E2430"


def _head(ws, title, sub):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13, color=INK)
    ws["A2"] = sub
    ws["A2"].font = Font(italic=True, size=9, color="555555")


def _table(ws, row, header, rows, widths, formats):
    thin = Side(style="thin", color="B9C0CC")
    for j, h in enumerate(header, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(bottom=thin)
    for i, r in enumerate(rows, row + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            if formats[j - 1]:
                c.number_format = formats[j - 1]
                c.alignment = Alignment(horizontal="right", vertical="top")
            else:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(bottom=thin)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def workbook(path, A):
    wb = Workbook()
    ws = wb.active
    ws.title = "Estimates"
    north = sorted({r for r, _ in A["warned"]})
    _head(ws, "%s of the %d estimates carry the low-reliability warning, all of them in the %s northern regions"
          % (WU.WORDS[A["n_warned"]], A["n_estimates"], WU.WORDS[len(north)].lower()),
          "Household Resources Survey 2027 tables. E: coefficient of variation %.1f per cent or more (QS-7 section 3)."
          % P.LINE)
    rows = []
    for r, name in WU.REGIONS:
        for y, label in WU.INDS:
            rows.append([r, name, label, A["est"][r][y], A["cv"][r][y], "E" if (r, y) in A["warned_set"] else "",
                         A["n_region"][r]])
    _table(ws, 4, ["Health region", "Region name", "Estimate", "Adults (thousands)", "Coefficient of variation (%)",
                   "Warning", "Responding adults in region"], rows, [9, 14, 42, 12, 14, 9, 13],
           [None, None, None, "0.0", "0.0", None, "#,##0"])
    for i, row in enumerate(rows, 5):
        if row[5]:
            for j in range(1, 8):
                ws.cell(row=i, column=j).font = Font(bold=True, color=RUST)
    ws.auto_filter.ref = "A4:G%d" % (4 + len(rows))
    ws.cell(row=6 + len(rows), column=1, value="Source: Household Resources Survey 2027, %s for the %s; analysis files "
            "frozen with the weights on 12 November 2027." % (P.AGENCY, P.CLIENT)).font = Font(italic=True, size=9,
                                                                                         color="555555")
    w2 = wb.create_sheet("Fieldwork")
    _head(w2, "Response ran from %.1f to %.1f per cent across the twelve regions (%.1f per cent overall)"
          % (A["rr_min"], A["rr_max"], A["rr_all"]),
          "Response rate as QS-7 section 4.2 defines it; calls as section 4.3 defines them, to the addresses with a full "
          "interview.")
    rows = [[r, name, A["issued_by_region"][r], A["interviews_by_region"][r], A["rr"][r] / 100.0, A["calls"][r]]
            for r, name in WU.REGIONS]
    rows.append(["All", "All regions", sum(A["issued_by_region"].values()), sum(A["interviews_by_region"].values()),
                 A["rr_all"] / 100.0, A["calls_all"]])
    _table(w2, 4, ["Health region", "Region name", "Addresses issued", "Full interviews", "Response rate",
                   "Calls to addresses with a full interview"], rows, [9, 14, 11, 11, 11, 18],
           [None, None, "#,##0", "#,##0", "0.0%", "#,##0"])
    for j in range(1, 7):
        w2.cell(row=4 + len(rows), column=j).font = Font(bold=True)
    w3 = wb.create_sheet("Chart")
    _head(w3, "Only the northern regions' care, arrears and North West's renting estimates cross the line",
          "Coefficient of variation by health region and estimate, per cent; the dashed line is the warning line, "
          "%.1f per cent." % P.LINE)
    w3.append([])
    w3.append(["Health region"] + [label for _, label in WU.INDS] + ["Warning line (%.1f%%)" % P.LINE])
    for r, name in WU.REGIONS:
        w3.append([name] + [A["cv"][r][y] for y, _ in WU.INDS] + [P.LINE])
    n = len(WU.REGIONS)
    bar = BarChart()
    bar.type = "col"
    bar.grouping = "clustered"
    bar.title = "Coefficient of variation by estimate (%)"
    bar.y_axis.title = "Coefficient of variation (%)"
    bar.y_axis.scaling.min = 0
    bar.y_axis.scaling.max = 50
    bar.y_axis.majorUnit = 10
    bar.y_axis.number_format = "0"
    bar.gapWidth = 60
    data = Reference(w3, min_col=2, max_col=5, min_row=4, max_row=4 + n)
    cats = Reference(w3, min_col=1, min_row=5, max_row=4 + n)
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(cats)
    for ser, colour in zip(bar.series, ("9AA5B8", "4F7CAC", RUST, "D9A441")):
        ser.graphicalProperties.solidFill = colour
        ser.graphicalProperties.line.solidFill = colour
    line = LineChart()
    line.add_data(Reference(w3, min_col=6, max_col=6, min_row=4, max_row=4 + n), titles_from_data=True)
    line.series[0].graphicalProperties.line.solidFill = INK
    line.series[0].graphicalProperties.line.dashStyle = "dash"
    line.series[0].graphicalProperties.line.width = 22000
    line.series[0].smooth = False
    bar += line
    bar.legend.position = "b"
    bar.height, bar.width = 10.5, 26
    w3.add_chart(bar, "A18")
    for j, wd in enumerate([14, 17, 17, 17, 19, 13], 1):
        w3.column_dimensions[get_column_letter(j)].width = wd
    for c in w3[4]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
    for row in w3.iter_rows(min_row=5, max_row=4 + n, min_col=2, max_col=6):
        for c in row:
            c.number_format = "0.0"
    w3.freeze_panes = "A5"
    w4 = wb.create_sheet("Notes")
    notes = [
        "Standard: QS-7 Statistical quality standard, contracted household surveys, edition 3 (sections 2 to 5).",
        "Estimates: adults aged 16 and over, in thousands, made with the final weights frozen on 12 November 2027, the "
        "design weights post-stratified to the Branch's population estimates for 1 July 2027 within the %d cells of "
        "health region, age group and sex. Don't know and refused are not counted; proxy answers are." % A["cells"],
        "Coefficient of variation: the standard error of the estimate as published over the unrounded estimate. The "
        "estimate is a post-stratified total, so its Taylor linearisation is each adult's departure from the weighted "
        "rate of their health region, age group and sex cell, weighted and totalled by PSU, the PSUs treated as selected "
        "with replacement within their strata (%d PSUs in %d strata). Held fixed, the final weights would put %d "
        "estimates over the line." % (A["psus"], A["strata"], A["n_fixed"]),
        "Health regions: the twelve regions as drawn by the Branch's 2025 boundary review. The master sample still holds "
        "the earlier region code for the %d PSUs the review moved; their addresses count in the regions they now lie "
        "in, where the analysis files place their households." % A["moved_psus"],
        "Full interviews: every household in the analysis files (%s). %d of them reached the office after the outcome "
        "extract of 30 September, which still shows them as non-contacts, other non-response or occupancy not "
        "established." % (WU.fmt(A["households"]), A["late"]),
        "Addresses never worked: %d issued addresses have no outcome; their eligibility was never established, and "
        "they count at the wave's eligible share (e = %.3f)." % (A["unworked"], A["e"]),
        "Calls: visits and telephone calls by the interviewer working each issue, in every issue. %d calls reached the "
        "office twice after a tablet's sync failed (the same address, issue, time, result and interviewer under a new "
        "record number); each counts once." % A["copies"],
        "Data as of 3 December 2027. Prepared by %s for the %s." % (P.AGENCY, P.CLIENT),
    ]
    w4["A1"] = "Notes"
    w4["A1"].font = Font(bold=True, size=12)
    for i, t in enumerate(notes, 3):
        w4.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True, vertical="top")
    w4.column_dimensions["A"].width = 110
    for sh in (ws, w2):
        sh.sheet_view.showGridLines = False
        sh.print_title_rows = "4:4"
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
        sh.sheet_properties.pageSetUpPr.fitToPage = True
    w3.page_setup.orientation = "landscape"
    w3.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)
    postprocess.scrub(path, P.AGENCY, "2027-12-03")
    postprocess.finish_office(path, "2027-12-02T15:10:00Z", "2027-12-03T19:25:00Z", P.AGENCY)


def note(path, A):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    doc = docx.Document()
    stl = doc.styles["Normal"]
    stl.font.name = "Calibri"
    stl.font.size = Pt(10.5)
    foot = doc.sections[0].footer.paragraphs[0]
    foot.text = "%s | Household Resources Survey 2027 | Quality note" % P.CLIENT
    foot.alignment = WD_ALIGN_PARAGRAPH.LEFT
    foot.runs[0].font.size = Pt(8)
    h = doc.add_paragraph()
    r = h.add_run("Household Resources Survey 2027: %s estimates carry the low-reliability warning" %
                  WU.WORDS[A["n_warned"]].lower())
    r.bold = True
    r.font.size = Pt(13)
    sub = doc.add_paragraph("Quality note published by the %s with the 2027 tables. Prepared by %s, 3 December 2027."
                            % (P.CLIENT, P.AGENCY))
    sub.runs[0].italic = True
    doc.add_paragraph(
        "The 2027 tables give %d estimates, four for each of the twelve health regions. %s of them carry the "
        "low-reliability warning, marked E in the tables: their coefficient of variation is %.1f per cent or more, so "
        "they should be used with caution. All %s are in the four northern regions, where the survey interviewed the "
        "fewest adults." % (A["n_estimates"], WU.WORDS[A["n_warned"]], P.LINE, WU.WORDS[A["n_warned"]].lower()))
    tb = doc.add_table(rows=1, cols=3)
    tb.style = "Light Grid Accent 1"
    for j, t in enumerate(("Health region", "Estimate", "Coefficient of variation (%)")):
        tb.rows[0].cells[j].text = t
    for r, y in A["warned"]:
        row = tb.add_row().cells
        row[0].text = "%s (%s)" % (WU.NAME[r], r)
        row[1].text = WU.IND_NAME[y]
        row[2].text = "%.1f" % A["cv"][r][y]
    src = doc.add_paragraph("Source: Household Resources Survey 2027. Coefficients of variation are computed before the "
                            "estimates are rounded for publication.")
    src.runs[0].font.size = Pt(8.5)
    src.runs[0].italic = True
    doc.add_paragraph(
        "Every other estimate stands without the warning, the public pension estimates of all twelve regions among "
        "them: in the four northern regions their coefficients of variation run from %.1f to %.1f per cent, and the "
        "highest coefficient of any unwarned estimate is %.1f per cent (%s in %s)."
        % (A["pen_north_min"], A["pen_north_max"], A["unwarned_max"], WU.IND_NAME[A["unwarned_max_key"][1]].lower(),
           WU.NAME[A["unwarned_max_key"][0]]))
    doc.add_paragraph(
        "How the warning is set. The estimates are made with weights scaled to the Branch's population estimates by "
        "health region, age group and sex. Each estimate's standard error is that of the estimate as published, by "
        "Taylor linearisation over the survey's strata and primary sampling units (QS-7 section 3). For an estimate "
        "made with weights scaled this way, the standard error rests on each adult's departure from the rate of their "
        "region, age group and sex, so an estimate that age and sex largely decide, such as the number of adults "
        "receiving a public pension, is far more precise than the spread of the weights alone would suggest.")
    doc.add_paragraph("Response. Response rates in the twelve regions ran from %.1f to %.1f per cent, %.1f per cent "
                      "overall (QS-7 section 4.2)." % (A["rr_min"], A["rr_max"], A["rr_all"]))
    s = doc.add_paragraph("Kristina Stephens, Lead Survey Statistician, %s" % P.AGENCY)
    s.paragraph_format.space_before = Pt(8)
    doc.core_properties.author = "Kristina Stephens"
    doc.core_properties.title = "Household Resources Survey 2027: quality note"
    doc.save(path)
    postprocess.scrub(path, "Kristina Stephens", "2027-12-03")
    postprocess.finish_office(path, "2027-12-03T14:30:00Z", "2027-12-03T20:05:00Z", P.AGENCY)


def write(folder, A):
    os.makedirs(folder, exist_ok=True)
    workbook(os.path.join(folder, XLSX), A)
    note(os.path.join(folder, DOCX), A)
