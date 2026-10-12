"""Assertions on the built pack: the ladder by rung and by grid cell, SF1 to SF15, the form 1 separation, the
rigorous family, the clean-data and lens-swap tests, the asks and the pair simulation, the input gates and the
realism sweep. run() returns the results and raises on any failure."""
import datetime as dt
import glob
import json
import os
import re
import zipfile

import numpy as np
import pandas as pd

import analysis as AN
import asks as AK
import pack as PK
import variants as VA
import world as W

GOLDEN_SET = ["N4", "D2", "D3", "D4", "D6"]
SETS = {"R0": ["N1", "N2", "N3", "N4", "N5"], "R1": ["N1", "N2", "N3", "N4", "N5", "D5"],
        "R2": ["N1", "N2", "N3", "N4", "N5", "D2", "D3", "D4", "D5", "D6"], "R3": ["N4", "D2", "D3", "D4"], "R4": GOLDEN_SET}
BANNED = re.compile(r"\b(pool|pooled|prevalence|vigilance|criterion|criteria)\b", re.I)


class Rec:
    def __init__(self):
        self.rows = []

    def __call__(self, name, ok, detail=""):
        self.rows.append((name, bool(ok), str(detail)))

    def failures(self):
        return [r for r in self.rows if not r[1]]


def _text_of(path):
    if path.endswith(".docx"):
        from docx import Document
        d = Document(path)
        t = "\n".join(p.text for p in d.paragraphs)
        for tb in d.tables:
            for row in tb.rows:
                t += "\n" + " ".join(c.text for c in row.cells)
        return t
    if path.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join(p.extract_text() for p in PdfReader(path).pages)
    if path.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(ws.title)
            for row in ws.iter_rows(values_only=True):
                out.extend(str(x) for x in row if isinstance(x, str))
        return "\n".join(out)
    if path.endswith(".parquet"):
        import pyarrow.parquet as pq
        return " ".join(pq.read_schema(path).names)
    if path.endswith((".csv", ".md", ".json")):
        return open(path, encoding="utf-8").read()
    return ""


def sets_equal(res, limits, want):
    return AN.designate(res, limits) == want


def run(target, quiet=True, strict=True):
    ck = Rec()
    pk = AN.Pack(target)
    out = VA.run(pk)
    L = out["limits"]
    R = out["R"]
    Qr, sc = out["Qr"], out["sc"]

    # ------------------------------------------------------------------ the ladder
    named = {"R0": R["R0"], "R1": R["R1"], "R2": R["R2"], "R3": R["R3"], "R4": R["golden"]}
    for k, r in named.items():
        ck(f"ladder {k} set", AN.designate(r, L) == SETS[k], AN.designate(r, L))
        need = 0.10 if k == "R4" else 0.05
        ck(f"ladder {k} margin", AN.margin(r, L) >= need, round(AN.margin(r, L), 3))
    ck("golden designation set", AN.designate(R["golden"], L) == GOLDEN_SET, AN.designate(R["golden"], L))
    ck("default path picks the best-fitting attribute model", out["best"] == min(out["aic"], key=out["aic"].get), out["best"])
    for name, r in out["att"].items():
        ck(f"attribute family files R1: {name}", AN.designate(r, L) == SETS["R1"], AN.designate(r, L))
    grid_sets = {k: tuple(AN.designate(r, L)) for k, r in out["grid"].items()}
    ck("grid: the roster toggle never moves a set", all(grid_sets[(0,) + k[1:]] == grid_sets[(1,) + k[1:]] for k in grid_sets), "")
    ck("grid: the answer only with both effects and the new client inside", [k for k, v in grid_sets.items() if list(v) == GOLDEN_SET] == [(0, 1, 1, 1), (1, 1, 1, 1)], "")
    ck("grid: six distinct sets", len(set(grid_sets.values())) == 6, sorted(set(grid_sets.values())))
    ck("grid: every cell feasible", all(len(v) < len(W.ORDER) for v in grid_sets.values()), "")

    # ------------------------------------------------------------------ the rigorous family
    for name, r in out["V"].items():
        ck(f"rigorous variant on the answer: {name}", AN.designate(r, L) == GOLDEN_SET, AN.designate(r, L))
        ck(f"rigorous variant margin 0.10: {name}", AN.margin(r, L) >= 0.10, round(AN.margin(r, L), 3))

    # ------------------------------------------------------------------ form 1 separation
    b, se = out["slope"]
    for q, (bq, sq) in out["slopes_within"].items():
        ck(f"form 1: rate slope clear of zero in queue {q}", bq + 1.96 * sq < 0, f"{bq:.3f} se {sq:.3f}")
    (bd, sd), (bn, sn) = out["slopes_within"]["D"], out["slopes_within"]["N"]
    ck("form 1: one slope on both queues", abs(bd - bn) < 1.96 * (sd ** 2 + sn ** 2) ** 0.5 and abs(bd - bn) < 1.96 * max(sd, sn) * 1.6, f"{bd:.3f} vs {bn:.3f}")
    for fam, (bf, sf) in out["fam_slopes"].items():
        ck(f"form 1: one slope across families ({fam})", abs(bf - b) < 1.96 * (sf ** 2 + se ** 2) ** 0.5, f"{bf:.3f} se {sf:.3f}")
    g_ho = out["heldout"]["golden"]
    for name, v in out["heldout"].items():
        if name == "golden":
            continue
        ck(f"form 1: held-out shift-days 15% worse: {name}", v >= 1.15 * g_ho, round(v / g_ho, 3))
    ck("form 1: shares add nothing once the rate is in", out["mix"]["p"] > 0.05, out["mix"])
    ck("form 1: November mix outside every closed shift-day", out["nov_share"]["D6"] > 3 * out["q3_max_share"]["D6"] and out["q3_max_share"]["D5"] == 0.0,
       (out["nov_share"], out["q3_max_share"]))
    ck("support: ten shift-days at or below November's day rate", out["support"]["D"] >= 10, out["support"])
    ck("support: ten shift-days at or above November's night rate", out["support"]["N"] >= 10, out["support"])
    ck("support: every roster verifier of the closed quarter worked both queues", out["verifiers_both_queues"], "")
    for nm, v in out["collapse"].items():
        if "hol" in nm:
            continue
        ck(f"queue and Saturday term collapses once the rate is in: {nm}", abs(v["z_with"]) < 2.0 and (abs(v["with_rate"]) < abs(v["without"]) or abs(v["z_without"]) < 2.0), v)
    for (q, m), v in out["monthly_pool"].items():
        lo, hi = (0.042, 0.044) if q == "D" else (0.024, 0.026)
        ck(f"monthly OCR error rate band {q} {m}", lo <= v <= hi, round(v * 100, 3))
    ck("default path's month-by-month back-test clean (every stream-month within 3 overdispersed standard errors)", out["backtest"][out["best"]] < 3.0, round(out["backtest"][out["best"]], 2))

    # ------------------------------------------------------------------ clean-data and lens swap
    for nm, d in out["T"].items():
        ck(f"clean-data: {nm}: golden unchanged", AN.designate(d["golden"], L) == GOLDEN_SET, AN.designate(d["golden"], L))
        ck(f"clean-data: {nm}: default unchanged", AN.designate(d["default"], L) == SETS["R1"], AN.designate(d["default"], L))
    gm = out["golden_model"]
    bm = AN.fit_spec(Qr, [a for a in AN.ATTRIBUTE if a["name"] == out["best"]][0])
    import glm
    C = Qr.C
    pg = glm.predict(gm, glm.design(C, AN._terms(AN.GOLDEN), gm["lv"])[0]) * C.n
    pb = glm.predict(bm, glm.design(C, AN._terms(bm["spec"]), bm["lv"])[0]) * C.n
    t = C.assign(pg=pg, pb=pb).groupby(["code", "month"])[["pg", "pb"]].sum()
    gap = float((abs(t.pg - t.pb) / t.pg).max())
    ck("lens swap: both models reproduce the closed quarter by stream-month (within 6%)", gap < 0.06, round(gap, 4))
    ck("lens swap: one November, two forecasts, two sets", AN.designate(R["golden"], L) != AN.designate(R["R1"], L), "")
    ck("own-stream-rate model files a fixed-detection set", AN.designate(out["own"], L) != GOLDEN_SET, AN.designate(out["own"], L))

    # ------------------------------------------------------------------ semantic forks
    qcp = _text_of(os.path.join(target, PK.FILES["qcp"]))
    ck("SF1 the whole month (4.2 'for the month')", "on second-pass verification for the month" in qcp, "")
    ck("SF2 value after first pass (4.1)", "value after first pass differs from the image" in qcp, "")
    ck("SF3 the plan's routing (4.1, 4.3) and R0 differs", "operating plan of record" in qcp and AN.designate(R["R0"], L) != GOLDEN_SET, "")
    ck("SF4 no forecast within 0.10 of its limit", AN.margin(R["golden"], L) >= 0.10, round(AN.margin(R["golden"], L), 3))
    sch = _text_of(os.path.join(target, PK.FILES["schedules"]))
    pairs = re.findall(r"([A-Z]{3}-[A-Z0-9]{3,4})\s+(\d+\.\d+)\s+(\d+\.\d+)%", sch)
    ck("SF5 limit and accuracy print one number", len(pairs) == 11 and all(abs(100 - float(l) / 10 - float(a)) < 1e-9 for _, l, a in pairs), len(pairs))
    kf = pk.kfields
    keyset = set(zip(kf.Stream, kf["Form version"], kf.Field))
    ck("SF6 the audit re-keys key fields only", set(zip(pk.au.stream_code, pk.au.form_version, pk.au.field_name)) <= keyset, "")
    ck("SF7 shift-days, not calendar days: same call", AN.designate(sf7_calendar(Qr, sc, gm), L) == GOLDEN_SET, "")
    ck("SF8 eleven streams", len(pk.streams.stream_code.unique()) == 11 and len(sc.plan.code.unique()) == 11, "")
    ck("SF9 the closed quarter is July to September", Qr.scr.day.max() <= dt.date(2026, 9, 30) and Qr.scr.day.min() >= dt.date(2026, 7, 1), "")
    ck("SF10 monthly-average November rates: same call", AN.designate(out["V"]["November rates as monthly averages"], L) == GOLDEN_SET, "")
    ck("SF11 misses from field edits plus the audit: same call", AN.designate(out["V"]["rate from field edits plus audited misses"], L) == GOLDEN_SET, "")
    ck("SF12 the closed quarter's roster: same call", AN.designate(out["T"]["closed-quarter roster"]["golden"], L) == GOLDEN_SET, "")
    ck("SF13 Hollings & Pryce routed from 1 November", "D6" in set(sc.plan.code) and sc.plan[sc.plan.code == "D6"].kf.sum() > 0, "")
    ck("SF14 verifier-introduced errors dropped: same call", AN.designate(out["T"]["verifier-introduced errors dropped"]["golden"], L) == GOLDEN_SET, "")
    ck("SF15 the divisor is key fields first-pass verified", "per 1,000 key fields first-pass verified" in qcp, "")

    # ------------------------------------------------------------------ asks and the pair simulation
    A = AK.run(target, pk, sc)
    for ask, stops in (("A", A["A"]), ("B", A["B"]), ("C", A["C"])):
        g = stops["golden"]
        for nm, v in stops.items():
            if nm == "golden":
                continue
            moved = [c for c in g if v[c] != g[c]]
            ck(f"ask {ask} stop {nm} differs from the golden", len(moved) >= 1, f"{len(moved)} of {len(g)} figures")
    ck("ask A: the cracker's dedupe misses every stream", all(A["A"]["dedupe"][c] != A["A"]["golden"][c] for c in W.ORDER), "")
    ck("ask A: summing every line misses every stream", all(A["A"]["sum"][c] != A["A"]["golden"][c] for c in W.ORDER), "")
    ck("ask B: the cracker misses every stream", all(A["B"]["cracker"][c] != A["B"]["golden"][c] for c in W.ORDER), "")
    ck("ask C: the cracker's clock misses every stream", all(A["C"]["utc_as_local"][c] != A["C"]["golden"][c] for c in AK.ten_streams()), "")
    for c in W.ORDER:
        frac = A["A_unrounded"][c] % 1
        ck(f"ask A {c} clear of the rounding edge", abs(frac - 0.5) > 0.05, round(frac, 3))
    for c in W.ORDER:
        frac = A["B_unrounded"][c] % 1
        ck(f"ask B {c} clear of the rounding edge", abs(frac - 0.5) > 0.02, round(frac, 3))
    for c in AK.ten_streams():
        u = A["C_unrounded"][c] * 10 % 1
        ck(f"ask C {c} clear of the rounding edge", abs(u - 0.5) > 0.1, round(A["C_unrounded"][c], 4))
    ck("pair simulation: every cracker row missed", all(r[2] for r in A["cracker_rows"]), [r for r in A["cracker_rows"] if not r[2]])
    ck("pair ceiling at or under 40", A["pair_ceiling_allowance"] <= 40, A["pair_ceiling_allowance"])
    ck("chart leader the same at nominal and budget rates", A["order"][0] == A["order_nominal"][0], (A["order"][0], A["order_nominal"][0]))

    # ------------------------------------------------------------------ input gates and realism
    files = sorted(os.listdir(target))
    ck("input gate: 10 or more files", len(files) >= 10, len(files))
    fmts = {f.rsplit(".", 1)[1] for f in files}
    ck("input gate: 3 or more formats", len(fmts) >= 3, sorted(fmts))
    ck("input gate: a file of 25,000 rows or more", len(pk.sp) >= 25000, len(pk.sp))
    ck("input gate: two distractors, unread by the solution", len(PK.DISTRACTORS) >= 2 and distractors_unread(), [PK.FILES[d] for d in PK.DISTRACTORS])
    banned_hits, dash_hits = [], []
    for f in files:
        txt = _text_of(os.path.join(target, f))
        if BANNED.search(txt):
            banned_hits.append((f, BANNED.search(txt).group(0)))
        if "—" in txt:
            dash_hits.append(f)
        if f.endswith((".csv",)):
            hdr = open(os.path.join(target, f)).readline()
            if BANNED.search(hdr):
                banned_hits.append((f, "header"))
    ck("no pack sentence or column carries the banned words", not banned_hits, banned_hits)
    ck("no em dash in the pack", not dash_hits, dash_hits)
    meth = re.compile(r"(logistic|regression|back-test|backtest|leave-one-out|model selection|best[- ]fitting|forecast method)", re.I)
    mhits = [f for f in files if f.endswith((".docx", ".pdf", ".md")) and meth.search(_text_of(os.path.join(target, f)))]
    ck("no pack document names, selects or scores a forecast method", not mhits, mhits)
    v = Qr.scr.groupby(["code", "day"]).key_fields.sum()
    dup = v.drop(index="D6", level=0).groupby(level=0).apply(lambda s: s.duplicated().mean())
    ck("realism: few repeated stream-day key-field totals", dup.max() < 0.15, dup.round(3).to_dict())
    ck("realism: no round headline totals", all(x % 1000 for x in [len(pk.sp), len(pk.ed), len(pk.au)]), "")
    ck("realism: no constant daily stream", v.groupby(level=0).std().min() > 0, "")
    mx = pd.to_datetime(pk.sp.closed_at).max()
    ck("realism: no record after the setting", mx < pd.Timestamp(W.ASOF), str(mx))
    office = [f for f in files if f.endswith((".xlsx", ".docx", ".pdf", ".pptx"))]
    import postprocess as PP
    s = PP._scrubber()
    bad = []
    for f in office:
        names, b2 = s.audit(os.path.join(target, f), PP.FLOOR, PP.CEILING)
        if names or b2:
            bad.append(f)
    ck("realism: every container scrubbed", not bad, bad)
    fails = ck.failures()
    out["asks"] = A
    out["checks"] = ck.rows
    if fails and strict:
        raise AssertionError("\n".join(f"FAIL {n}: {d}" for n, _, d in fails))
    return out


def distractors_unread():
    here = os.path.dirname(os.path.abspath(__file__))
    src = "".join(open(os.path.join(here, f)).read() for f in ("analysis.py", "asks.py", "variants.py"))
    return all(PK.FILES[d] not in src for d in PK.DISTRACTORS)


def sf7_calendar(Qr, sc, gm):
    """November's night rates read by calendar day (half of each night shift on the next date): same call"""
    pl = sc.plan.copy()
    n = pl[pl.queue == "N"]
    nxt = n.assign(day=n.day + dt.timedelta(days=1), kf=n.kf / 2)
    pl2 = pd.concat([pl[pl.queue == "D"], n.assign(kf=n.kf / 2), nxt])
    pl2 = pl2[pl2.day <= W.NOV[1]]
    orig = sc.plan
    sc.plan = pl2.groupby(["code", "day", "queue", "kfpd"], as_index=False).kf.sum()
    sc.plan["docs"] = sc.plan.kf / sc.plan.kfpd
    try:
        return AN.forecast(Qr, sc, gm)
    finally:
        sc.plan = orig
