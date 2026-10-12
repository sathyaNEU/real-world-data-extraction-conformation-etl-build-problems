"""Everything computed from the shipped files under target/: the closed quarter's audited detection, the
November forecasts under every candidate method (the ladder's rungs, the attribute family, the rigorous
variants), the asks and the cracker's stops. checks.py asserts on it; golden.py writes the deliverables from it."""
import datetime as dt
import json
import math
import os
import re

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.stats import chi2

import glm
import world as W

DAY = dt.timedelta(days=1)
SID2 = {v["sid"]: k for k, v in W.STREAMS.items()}
FAM = {k: v["family"] for k, v in W.STREAMS.items()}


# ----------------------------------------------------------------------------------------------- loading
def _xl(path, sheet, header=0):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    h = rows[header]
    return pd.DataFrame([r for r in rows[header + 1:] if any(x is not None for x in r)], columns=h)


def limits_from_pdf(path):
    from pypdf import PdfReader
    txt = "\n".join(p.extract_text() for p in PdfReader(path).pages)
    out = {}
    for sid, lim in re.findall(r"([A-Z]{3}-[A-Z0-9]{3,4})\s+(\d+\.\d+)\s+\d+\.\d+%", txt):
        out[SID2[sid]] = float(lim)
    return out


def credit_pct_from_pdf(path):
    from pypdf import PdfReader
    txt = "\n".join(p.extract_text() for p in PdfReader(path).pages)
    out = {}
    for sid, pct in re.findall(r"([A-Z]{3}-[A-Z0-9]{3,4})\s+\d+\.\d+\s+\d+\.\d+%\s+(\d+(?:\.\d+)?)%", txt):
        out[SID2[sid]] = float(pct)
    return out


class Pack:
    def __init__(self, target):
        self.t = target
        f = lambda n: os.path.join(target, n)
        self.sp = pd.read_parquet(f("first_pass_verifications_2026q3.parquet"))
        self.ed = pd.read_parquet(f("first_pass_field_edits_2026q3.parquet"))
        self.au = pd.read_parquet(f("blind_rekey_audit_2026q3.parquet"))
        self.emp = _xl(f("workforce_register.xlsx"), "Employees")
        self.badge = _xl(f("workforce_register.xlsx"), "Badge allocation")
        self.streams = pd.read_csv(f("stream_register.csv"))
        self.vers = _xl(f("form_specifications.xlsx"), "Versions")
        self.kfields = _xl(f("form_specifications.xlsx"), "Key fields")
        self.plan_daily = _xl(f("november_operating_plan.xlsx"), "Daily key fields")
        self.roster = _xl(f("november_operating_plan.xlsx"), "Roster")
        self.routing = _xl(f("november_operating_plan.xlsx"), "Routing")
        br = _xl(f("november_operating_plan.xlsx"), "Budget rates")
        self.budget = {r[0]: float(r[1]) for r in br.itertuples(index=False) if r[0] in ("USD", "GBP", "AUD")}
        self.limits = limits_from_pdf(f("client_quality_schedules_2026.pdf"))
        self.credit = credit_pct_from_pdf(f("client_quality_schedules_2026.pdf"))


# ----------------------------------------------------------------------------------------------- the closed quarter
def kfpd_table(pk):
    v = pk.vers.rename(columns={"Stream": "sid", "Form version": "version", "Effective from (shift-day)": "eff",
                                "Key fields per document": "kfpd"})
    v["code"] = v.sid.map(SID2)
    v["eff"] = pd.to_datetime(v.eff).dt.date
    return v


def kfpd_on(vt, code, d):
    g = vt[(vt.code == code) & (vt.eff <= d)].sort_values("eff")
    return int(g.kfpd.iloc[-1])


def screens(pk, keep_timeouts=False, naive_badge=False, calendar_day=False):
    sp = pk.sp if keep_timeouts else pk.sp[pk.sp.outcome == "SUBMITTED"]
    sp = sp.copy()
    sp["code"] = sp.stream_code.map(SID2)
    qmap = {"N1": "N", "N2": "N", "N3": "N", "N4": "N", "N5": "N", "D1": "D", "D2": "D", "D3": "D", "D4": "D", "D5": "D", "D6": "D"}
    sp["queue"] = sp.code.map(qmap)
    od = sp.opened_at
    shift_day = od.dt.normalize() - pd.to_timedelta(((sp.queue == "N") & (od.dt.hour < 12)).astype(int), unit="D")
    sp["day"] = shift_day.dt.date
    start = pd.to_datetime(sp.day) + pd.to_timedelta(np.where(sp.queue == "N", 20.5, 8.0), unit="h")
    sp["hours"] = (sp.closed_at - start).dt.total_seconds() / 3600
    # badge to employee by date
    if naive_badge:
        sp["emp"] = sp.badge_no
    else:
        b = pk.badge.copy()
        b["Valid from"] = pd.to_datetime(b["Valid from"])
        b["Valid to"] = pd.to_datetime(b["Valid to"]).fillna(pd.Timestamp("2099-01-01"))
        sp["emp"] = None
        d = pd.to_datetime(sp.day)
        emp = np.empty(len(sp), dtype=object)
        for r in b.itertuples():
            m = (sp.badge_no.values == r.Badge) & (d.values >= np.datetime64(r[3])) & (d.values <= np.datetime64(r[4]))
            emp[m] = r[2]
        sp["emp"] = emp
    hire = dict(zip(pk.emp["Emp No"], pd.to_datetime(pk.emp["Date of joining"]).dt.date))
    if naive_badge:
        bh = pk.badge.sort_values("Valid from").groupby("Badge")["Emp No"].first()
        sp["tenure"] = [(d - hire[bh[b]]).days for d, b in zip(sp.day, sp.badge_no)]
    else:
        sp["tenure"] = [(d - hire[e]).days for d, e in zip(sp.day, sp.emp)]
    if calendar_day:
        sp["day"] = od.dt.normalize().dt.date
    wl = sp.groupby(["emp", "day"]).size().rename("wl")
    sp = sp.join(wl, on=["emp", "day"])
    return sp


def ten_band(days):
    days = np.asarray(days)
    return np.where(days < 90, "new", np.where(days < 180, "mid", "exp"))


def audit_frame(pk, scr):
    a = pk.au.copy()
    kf = pk.kfields.rename(columns={"Stream": "sid", "Form version": "version", "Field": "field_name", "Field type": "ftype"})
    a = a.merge(kf, left_on=["stream_code", "form_version", "field_name"], right_on=["sid", "version", "field_name"], how="left")
    cols = ["document_id", "day", "queue", "emp", "hours", "tenure", "wl", "code"]
    s = scr[cols]
    if s.document_id.duplicated().any():
        a = a.merge(s, on="document_id", how="inner")
    else:
        a = a.merge(s, on="document_id", how="left")
    a["err"] = a.result_code.isin(["C", "R"])
    a["miss"] = a.result_code.eq("R")
    a["fa"] = a.result_code.eq("V")
    return a


def pool_rates(a, scr, mode="strat", pk=None):
    """the OCR error rate of every key field a queue carried on a shift-day"""
    fields = scr.groupby(["day", "queue", "code"]).key_fields.sum().rename("N")
    if mode in ("strat", "plain"):
        st = a.groupby(["day", "queue", "code"]).agg(e=("err", "sum"), n=("err", "size")).join(fields)
        if mode == "plain":
            g = st.groupby(["day", "queue"])[["e", "n"]].sum()
            return (g.e / g.n).rename("pool")
        st["est"] = st.e / st.n * st.N
        g = st.groupby(["day", "queue"])[["est", "N"]].sum()
        return (g.est / g.N).rename("pool")
    if mode == "edit":
        ed = pk.ed.merge(scr[["screen_id", "day", "queue"]], on="screen_id", how="inner")
        edits = ed.groupby(["day", "queue"]).size().rename("k")
        tot = fields.groupby(["day", "queue"]).sum()
        aud = a.groupby(["day", "queue"]).agg(m=("miss", "sum"), e=("err", "sum"), v=("fa", "sum"), n=("err", "size"))
        fa_rate = aud.v / (aud.n - aud.e)
        mr = aud.m / aud.e
        det = edits - fa_rate * tot
        return (det / (1 - mr) / tot).rename("pool")
    raise ValueError(mode)


def holiday_class(day, queue):
    """weekday first, so a weekday is every model's reference level"""
    if day.weekday() == 5:
        return "2sat"
    if queue == "N" and day in W.US_HOLIDAYS_Q3 + W.US_HOLIDAYS_NOV:
        return "3hol"
    if queue == "D" and day in W.UK_HOLIDAYS_Q3:
        return "3hol"
    return "1wk"


class Quarter:
    """the audited OCR errors of the closed quarter as fitting cells, with every candidate covariate"""

    def __init__(self, pk, pool_mode="strat", keep_timeouts=False, naive_badge=False, calendar_day=False, drop_fa=False):
        self.pk = pk
        self.scr = screens(pk, keep_timeouts, naive_badge, calendar_day)
        self.a = audit_frame(pk, self.scr)
        if drop_fa:
            self.a = self.a.assign(fa=False)
        self.pool = pool_rates(self.a, self.scr, pool_mode, pk)
        a = self.a
        e = a[a.err].join(self.pool, on=["day", "queue"])
        e["fam"] = e.code.map(FAM)
        e["ten"] = ten_band(e.tenure.values)
        e["tos"] = np.where(e.hours > 5, "late", "early")
        e["ton"] = np.where((e.queue == "N") & (e.hours >= 5.5), "small_hours", "other")
        e["hcls"] = [holiday_class(d, q) for d, q in zip(e.day, e.queue)]
        e["qdow"] = e.queue + "_" + e.hcls
        q = np.quantile(e.wl, [1 / 3, 2 / 3])
        e["wlb"] = np.where(e.wl <= q[0], "low", np.where(e.wl <= q[1], "mid", "high"))
        e["lp"] = np.log(e.pool)
        e["month"] = pd.to_datetime(e.day).dt.month
        self.E = e
        self.keys = ["queue", "day", "code", "fam", "ftype", "emp", "ten", "tos", "ton", "qdow", "wlb", "month", "pool"]
        self.C = e.groupby(self.keys).agg(y=("miss", "sum"), n=("miss", "size")).reset_index()
        self.C["lp"] = np.log(self.C.pool)
        self.C["lin"] = self.C.pool
        self.C["lgt"] = np.log(self.C.pool / (1 - self.C.pool))
        self.fa_rate = a.fa.sum() / (~a.err).sum()
        ae = a.groupby(["code", "ftype"]).agg(e=("err", "sum"), n=("err", "size")).reset_index()
        self.ftmix = ae
        self.p_hat = (ae.groupby("code").e.sum() / ae.groupby("code").n.sum()).to_dict()
        self.emps = sorted(self.C.emp.unique())

    def q3_measured(self):
        a = self.a
        g = a.groupby("code").agg(r=("miss", "sum"), v=("fa", "sum"), n=("err", "size"))
        return ((g.r + g.v) / g.n * 1000).to_dict()

    def monthly_pool(self):
        f = self.scr.groupby(["day", "queue"]).key_fields.sum().rename("N")
        p = pd.concat([self.pool, f], axis=1).reset_index()
        p["month"] = pd.to_datetime(p.day).dt.month
        return {(q, m): float((g.pool * g.N).sum() / g.N.sum()) for (q, m), g in p.groupby(["queue", "month"])}


# ----------------------------------------------------------------------------------------------- models
GOLDEN = dict(name="golden", x="lp", terms=[("cat", "fam"), ("cat", "ftype"), ("cat", "emp"), ("cat", "ten")], link="logit")
ATTRIBUTE = [
    dict(name="A1 family, field type, queue", x=None, terms=[("cat", "fam"), ("cat", "ftype"), ("cat", "queue")], link="logit"),
    dict(name="A2 + verifier", x=None, terms=[("cat", "fam"), ("cat", "ftype"), ("cat", "queue"), ("cat", "emp")], link="logit"),
    dict(name="A3 + tenure", x=None, terms=[("cat", "fam"), ("cat", "ftype"), ("cat", "queue"), ("cat", "emp"), ("cat", "ten")], link="logit"),
    dict(name="A4 + queue by day type", x=None, terms=[("cat", "fam"), ("cat", "ftype"), ("cat", "qdow"), ("cat", "emp"), ("cat", "ten")], link="logit"),
    dict(name="A5 + time on shift", x=None, terms=[("cat", "fam"), ("cat", "ftype"), ("cat", "qdow"), ("cat", "emp"), ("cat", "ten"), ("cat", "tos")], link="logit"),
    dict(name="A6 + time of night, workload", x=None, terms=[("cat", "fam"), ("cat", "ftype"), ("cat", "qdow"), ("cat", "emp"), ("cat", "ten"), ("cat", "tos"), ("cat", "ton"), ("cat", "wlb")], link="logit"),
]
NUISANCE = ["tos", "ton", "wlb"]


def _terms(spec):
    t = list(spec["terms"])
    if spec.get("x"):
        t = [spec["x"]] + t
    return t


def fit_spec(Qr, spec, C=None, offset=None):
    C = Qr.C if C is None else C
    X, names, lv = glm.design(C, _terms(spec))
    off = None if offset is None else offset(C)
    m = glm.fit(X, C.y, C.n, link=spec.get("link", "logit"), offset=off)
    m.update(names=names, lv=lv, spec=spec)
    return m


class Scenario:
    """November as the plan of record files it, or a toggled version of it"""

    def __init__(self, pk, Qr):
        self.pk, self.Qr = pk, Qr
        vt = kfpd_table(pk)
        pdl = pk.plan_daily.copy()
        pdl["day"] = pd.to_datetime(pdl["Shift-day"]).dt.date
        rows = []
        for c in W.ORDER:
            for d, v in zip(pdl.day, pdl[W.SID[c]]):
                if v is not None and not pd.isna(v) and v > 0:
                    rows.append((c, d, float(v), kfpd_on(vt, c, d)))
        self.plan = pd.DataFrame(rows, columns=["code", "day", "kf", "kfpd"])
        self.plan["docs"] = self.plan.kf / self.plan.kfpd
        rt = pk.routing
        self.nov_queue = {SID2[s]: ("N" if q == "Night" else "D") for s, q in zip(rt.Stream, rt["Queue (November)"])}
        self.plan["queue"] = self.plan.code.map(self.nov_queue)
        ro = pk.roster
        self.roster = pd.DataFrame({"emp": ro["Emp No"], "queue": ro.Queue.map({"Day": "D", "Night": "N"})})
        hire = dict(zip(pk.emp["Emp No"], pd.to_datetime(pk.emp["Date of joining"]).dt.date))
        self.hire = hire

    def pools(self, p_hat):
        pl = self.plan.assign(e=self.plan.kf * self.plan.code.map(p_hat))
        g = pl.groupby(["queue", "day"])[["e", "kf"]].sum()
        return (g.e / g.kf).rename("pool")

    def monthly(self, p_hat):
        pl = self.plan.assign(e=self.plan.kf * self.plan.code.map(p_hat))
        g = pl.groupby("queue")[["e", "kf"]].sum()
        return (g.e / g.kf).to_dict()


def nuisance_weights(Qr, queue, col):
    E = Qr.E[Qr.E.queue == queue]
    w = E[col].value_counts(normalize=True)
    return w.to_dict()


def forecast(Qr, sc, m, night_fx=True, day_fx=True, hp_in_model=True, roster=True, monthly=False, xfun=None,
             p_hat=None, slope=None):
    """each stream's November first-pass residual per 1,000 key fields under fitted model m"""
    spec = m["spec"]
    lv = m["lv"]
    p_hat = p_hat or Qr.p_hat
    pools = sc.pools(p_hat)
    mon = sc.monthly(p_hat)
    q3 = {q: float((Qr.pool.xs(q, level="queue") * Qr.scr[Qr.scr.queue == q].groupby("day").key_fields.sum()).sum()
                   / Qr.scr[Qr.scr.queue == q].key_fields.sum()) for q in ("D", "N")}
    terms = _terms(spec)
    used = {t[1] for t in terms if isinstance(t, tuple)}
    emp_lv = lv.get("emp", [])
    emp_cols = [j for j, nm in enumerate(m["names"]) if nm.startswith("emp=")]
    rows = []
    for code in W.ORDER:
        q = sc.nov_queue[code]
        g = sc.plan[sc.plan.code == code]
        mix = Qr.ftmix[Qr.ftmix.code == code]
        crew = sc.roster[sc.roster.queue == q].emp.tolist() if roster else None
        for r in g.itertuples():
            if monthly:
                pool = mon[q]
            else:
                pool = float(pools[(q, r.day)])
            if (q == "N" and not night_fx) or (q == "D" and not day_fx) or (code == "D6" and not hp_in_model):
                pool = q3[q]
            hc = holiday_class(r.day, q)
            members = crew if roster else ["__q3__"]
            for e_ in members:
                if roster:
                    ten = ten_band([(r.day - sc.hire[e_]).days])[0]
                else:
                    ten = "__q3__"
                for ft in mix.itertuples():
                    rows.append(dict(code=code, day=r.day, kf=r.kf, queue=q, fam=FAM[code], ftype=ft.ftype, emp=e_, ten=ten,
                                     qdow=f"{q}_{hc}", pool=pool, wft=ft.n, pft=ft.e / ft.n, nm=len(members)))
    D = pd.DataFrame(rows)
    D["lp"] = np.log(D.pool)
    D["lin"] = D.pool
    D["lgt"] = np.log(D.pool / (1 - D.pool))
    if xfun is not None:
        D[spec["x"]] = xfun(D)
    # expand nuisance terms over their closed-quarter mix in the queue
    for col in NUISANCE:
        if col in used:
            parts = []
            for q in ("D", "N"):
                w = nuisance_weights(Qr, q, col)
                sub = D[D.queue == q]
                for k, v in w.items():
                    parts.append(sub.assign(**{col: k}, wn=v))
            D = pd.concat(parts, ignore_index=True)
            D["wq"] = D.get("wq", 1.0) * D.wn
            D = D.drop(columns="wn")
    if "wq" not in D:
        D["wq"] = 1.0
    X, _, _ = glm.design(D.assign(emp=D.emp.where(D.emp.isin(emp_lv), emp_lv[0] if emp_lv else "")), terms, lv)
    if emp_lv:
        unk = ~D.emp.isin(emp_lv).values
        if unk.any():
            X[np.ix_(unk, emp_cols)] = 1.0 / len(emp_lv)
    if not roster and "ten" in used:
        # the closed quarter's tenure mix in the queue
        cols = [j for j, nm in enumerate(m["names"]) if nm.startswith("ten=")]
        for q in ("D", "N"):
            w = nuisance_weights(Qr, q, "ten")
            sel = (D.queue == q).values
            for j in cols:
                lvl = m["names"][j].split("=", 1)[1]
                X[sel, j] = w.get(lvl, 0.0)
        if emp_lv:
            for q in ("D", "N"):
                E = Qr.E[Qr.E.queue == q]
                w = E.emp.value_counts(normalize=True)
                sel = (D.queue == q).values
                for j in emp_cols:
                    X[sel, j] = w.get(m["names"][j].split("=", 1)[1], 0.0)
    if slope is not None:
        j = m["names"].index(spec["x"])
        b = m["b"].copy()
        b[j] = slope
        mm = dict(m, b=b)
        pmiss = glm.predict(mm, X)
    else:
        pmiss = glm.predict(m, X)
    D["pm"] = pmiss
    D["contrib"] = D.pft * D.pm * D.wft * D.wq / D.nm
    D["wsum"] = D.wft * D.wq / D.nm
    per = D.groupby(["code", "day", "kf"])[["contrib", "wsum"]].sum().reset_index()
    per["r"] = per.contrib / per.wsum
    out = {}
    for code, g in per.groupby("code"):
        p = p_hat[code]
        r = (g.r * g.kf).sum() / g.kf.sum()
        out[code] = (r + Qr.fa_rate * (1 - p)) * 1000
    return out


def designate(res, limits):
    return sorted([c for c in W.ORDER if res[c] > limits[c]], key=W.ORDER.index)


def ratios(res, limits):
    return {c: res[c] / limits[c] for c in W.ORDER}


def margin(res, limits):
    return min(abs(res[c] / limits[c] - 1) for c in W.ORDER)
