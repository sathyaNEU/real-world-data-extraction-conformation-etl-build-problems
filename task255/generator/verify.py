"""The independent verifier: reads only the shipped bundle (target/) and the goldens, on its own code path (its own
parsing of the PDFs' figures, its own likelihood in the log of sigma, a quasi-Newton optimiser), and recomputes the
call, every attachment's expected count, each coverage group's severity, the premiums and the rated exposure by
class."""
import csv
import datetime as dt
import math
import os
import re
import sys
from collections import Counter, defaultdict

import numpy as np
import openpyxl
from scipy import optimize, special


def pdf_text(path):
    from pypdf import PdfReader
    return re.sub(r"\s+", " ", " ".join(p.extract_text() for p in PdfReader(path).pages))


def find(folder, prefix):
    hits = [f for f in os.listdir(folder) if f.startswith(prefix)]
    assert len(hits) == 1, (prefix, hits)
    return os.path.join(folder, hits[0])


def norm_sf_log(z):
    return np.log(0.5 * special.erfc(z / math.sqrt(2.0)))


def main(target, golden):
    memo = pdf_text(find(target, "liability_coverage_memorandum"))
    study = pdf_text(find(target, "actuarial_study"))
    quotes = pdf_text(find(target, "excess_liability_indications"))
    policy = pdf_text(find(target, "risk_financing_policy"))
    m = re.search(r"\$([\d,]+) for occurrences in (\d{4}) to (\d{4}) and \$([\d,]+) for occurrences from 1 January (\d{4})", memo)
    r_old, y0, y1, r_new, y2 = int(m.group(1).replace(",", "")), int(m.group(2)), int(m.group(3)), int(m.group(4).replace(",", "")), int(m.group(5))
    retention = lambda ay: r_old if ay <= y1 else r_new
    index = {int(a): float(b) for a, b in re.findall(r"\b(20\d\d)(?: \(projected\))? (\d{2,3}\.\d)\b", study)}
    lam = {"LE": float(re.search(r"Law enforcement liability \(police department payroll\) \$[\d,]+ [\d.]+ ([\d.]+)",
                                 study).group(1)),
           "OT": float(re.search(r"\(all other payroll\) \$[\d,]+ [\d.]+ ([\d.]+)", study).group(1))}
    LE = {"LEL", "LAWE"}
    win = re.search(r"accident years (\d{4}) to (\d{4})\)", policy)
    years = range(int(win.group(1)), int(win.group(2)) + 1)
    opts = [(int(a.replace(",", "")), float(gl), float(al)) for a, gl, al in
            re.findall(r"\$([\d,]{9,}) \$10,000,000 \$(\d\.\d{4}) \$([\d.]+)", quotes)]
    assert len(opts) == 7 and len(index) >= 14, (len(opts), len(index))
    line = 1.0 if "is one or fewer" in policy else None

    # the pool's record of every occurrence of the policy's years, both administrators, an open file at what the pool
    # has paid and still reserves
    occ = []
    wb = openpyxl.load_workbook(find(target, "liability_closed_claims"), read_only=True)
    for name in wb.sheetnames:
        if not name.startswith("AY"):
            continue
        for row in wb[name].iter_rows(min_row=6, values_only=True):
            if not row or not row[0]:
                continue
            ay = dt.datetime.strptime(row[3], "%m/%d/%Y").year
            occ.append((ay, float(row[7]), float(row[7]) + float(row[10]), row[6] == "CL", row[2] in LE))
    for r in csv.DictReader(open(find(target, "liability_loss_run"))):
        ay = int(r["loss_date"][:4])
        occ.append((ay, float(r["deductible"]), float(r["deductible"]) + float(r["paid_indemnity"]) +
                    float(r["paid_defense"]) + float(r["reserve_indemnity"]) + float(r["reserve_defense"]),
                    r["status"] == "Closed", r["coverage"] in LE))
    fit = [o for o in occ if o[0] in years]
    assert sum(1 for o in fit if not o[3]) == 8, "the policy's years hold eight open files"

    def fit_group(sel):
        f = np.array([index[2027] / index[o[0]] for o in sel])
        x = np.log(np.array([o[2] for o in sel]) * f)
        d = np.log(np.array([o[1] for o in sel]) * f)
        cens = np.array([o[2] >= retention(o[0]) - 0.005 for o in sel])

        def nll(p):
            mu, ls = p
            s = math.exp(ls)
            z = (x - mu) / s
            dens = -0.5 * z * z - 0.5 * math.log(2 * math.pi) - ls - x
            return -(np.where(cens, norm_sf_log(z), dens) - norm_sf_log((d - mu) / s)).sum()

        start = [float(x.mean()), math.log(float(x.std()))]
        res = optimize.minimize(nll, start, method="BFGS", options=dict(gtol=1e-9))
        res = optimize.minimize(nll, res.x, method="Nelder-Mead", options=dict(xatol=1e-12, fatol=1e-12, maxiter=20000))
        return float(res.x[0]), math.exp(float(res.x[1])), int(cens.sum()), len(sel)

    curves = {"LE": fit_group([o for o in fit if o[4]]), "OT": fit_group([o for o in fit if not o[4]])}

    def count(a):
        return sum(lam[g] * math.exp(norm_sf_log((math.log(a) - mu) / sigma) - norm_sf_log((math.log(25_000) - mu) / sigma))
                   for g, (mu, sigma, _, _) in curves.items())

    counts = {a: count(a) for a, _, _ in opts}
    limit = int(re.search(r"Each option provides \$([\d,]+) for each occurrence above its attachment", quotes)
                .group(1).replace(",", ""))

    def lev(mu, sigma, d):
        z = (math.log(d) - mu) / sigma
        return (math.exp(mu + sigma * sigma / 2) * 0.5 * special.erfc(-(z - sigma) / math.sqrt(2.0))
                + d * 0.5 * special.erfc(z / math.sqrt(2.0)))

    def layer(a):
        return sum(lam[g] * (lev(mu, sigma, a + limit) - lev(mu, sigma, a))
                   / math.exp(norm_sf_log((math.log(25_000) - mu) / sigma)) for g, (mu, sigma, _, _) in curves.items())

    layers = {a: int(round(layer(a) / 1e5) * 100_000) for a, _, _ in opts}
    call = next(a for a, _, _ in opts if counts[a] <= line)

    # the rated exposure: the 2027 members (a withdrawal effective by 31 December 2026 leaves a member out), each
    # member's last declaration once; one vehicle per latest unit number added for it, under the member on that
    # number's latest row in effect at the rating date
    leaving = set()
    for r in csv.DictReader(open(find(target, "member_roster"))):
        if r["withdrew"] and dt.date.fromisoformat(r["withdrew"]) <= dt.date(2026, 12, 31):
            leaving.add(int(r["member_no"]))
    decl = {}
    wb = openpyxl.load_workbook(find(target, "exposure_declarations"), read_only=True)
    cls = {}
    for row in wb["Declarations"].iter_rows(min_row=2, values_only=True):
        cls[row[0]] = row[1]
        if row[6] is None or row[0] in leaving:
            continue
        sub = row[7]
        if row[0] not in decl or sub > decl[row[0]][0]:
            decl[row[0]] = (sub, int(row[6]))
    toe = {k: v for k, (_, v) in decl.items()}
    rating = dt.date(2026, 9, 30)
    adds, gone, key, owner = {}, set(), {}, {}

    def ident(v, s):
        v = re.sub(r"[\s-]", "", v).upper()
        return v if len(v) == 17 else "S:" + re.sub(r"[\s-]", "", s).upper()

    log = sorted(csv.DictReader(open(find(target, "vehicle_schedule_log"))),
                 key=lambda r: (r["effective_date"], r["entered_date"]))
    for r in log:
        k = r["unit_no"]
        eff = dt.date.fromisoformat(r["effective_date"])
        if eff > rating:
            continue
        if r["txn_type"] == "ADD":
            adds[k] = (eff, r["entered_date"])
            key[k] = ident(r["vin"], r["serial_no"])
        elif r["txn_type"] == "DELETE":
            gone.add(k)
        owner[k] = int(r["member_no"])
    latest = {}
    for k, stamp in adds.items():
        if key[k] not in latest or stamp > latest[key[k]][0]:
            latest[key[k]] = (stamp, k)
    units = Counter(owner[k] for _, (_, k) in latest.items() if k not in gone and owner[k] not in leaving)
    T, U = sum(toe.values()), sum(units.values())
    prem = {a: round(gl * T / 1000 + al * U) for a, gl, al in opts}
    by = defaultdict(lambda: [0, 0, 0])
    for mno, v in toe.items():
        by[cls[mno]][0] += 1
        by[cls[mno]][1] += v
        by[cls[mno]][2] += units.get(mno, 0)

    # the goldens
    gw = openpyxl.load_workbook(os.path.join(golden, "retention_2027.xlsx"))
    ws = gw["Retention 2027"]
    rows = [r for r in ws.iter_rows(min_row=6, max_row=12, values_only=True)]
    for r in rows:
        assert abs(r[1] - round(counts[r[0]], 1)) < 1e-9, ("count", r[0], r[1], counts[r[0]])
        assert r[3] == layers[r[0]], ("layer loss", r[0], r[3], layers[r[0]])
        assert (r[2] == "Yes") == (counts[r[0]] <= line)
    rec_txt = ws["A2"].value
    assert "${:,.0f}".format(call) in rec_txt, rec_txt
    for g, lab in (("LE", "Severity, law enforcement"), ("OT", "Severity, general")):
        txt = [c for c in ws.iter_rows(min_row=13, values_only=True) if c[0] and str(c[0]).startswith(lab)][0][1]
        assert "log mean %.2f" % curves[g][0] in txt and "log standard deviation %.2f" % curves[g][1] in txt, txt
    print("verify: from target/ alone, %d occurrences; law enforcement %d (%d censored) log mean %.4f, log sd %.4f; other "
          "%d (%d censored) log mean %.4f, log sd %.4f; call $%s (%.2f; runner-up %.2f); premiums %s; layer losses %s; "
          "operating expenditures $%s, %d vehicles; goldens agree" % (
              len(fit), curves["LE"][3], curves["LE"][2], curves["LE"][0], curves["LE"][1], curves["OT"][3],
              curves["OT"][2], curves["OT"][0], curves["OT"][1], "{:,}".format(call), counts[call],
              counts[[a for a, _, _ in opts][[a for a, _, _ in opts].index(call) - 1]],
              ", ".join("{:,}".format(prem[a]) for a, _, _ in opts), ", ".join("{:,}".format(layers[a]) for a, _, _ in opts),
              "{:,}".format(T), U))


if __name__ == "__main__":
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "target"),
         sys.argv[2] if len(sys.argv) > 2 else os.path.join(here, "golden"))
