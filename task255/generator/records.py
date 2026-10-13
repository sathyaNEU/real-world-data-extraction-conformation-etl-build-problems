"""The pack's data files, each as its source system exports it: the legacy administrator's closed-claim archive, the
current administrator's loss run, the vehicle schedule log, the 2027 exposure declarations, the member roster, and
two files the pool keeps beside them (the deductible history and the actuary's triangles)."""
import csv
import datetime as dt
import hashlib

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

import world as W

LEGACY = "liability_closed_claims_tallowmark_2014-2020.xlsx"
CURRENT = "liability_loss_run_brisbrook_2026-09-30.csv"
LOG = "vehicle_schedule_log_2022-2026.csv"
DECL = "exposure_declarations_2027.xlsx"
ROSTER = "member_roster.csv"
DEDHIST = "member_deductible_history.csv"
TRIANGLES = "liability_triangles_2026-09-30.csv"

LEGACY_PART = {"GL": "GL", "AL": "AL", "LE": "LEL", "PO": "POL", "EP": "EPL"}
CURRENT_PART = {"GL": "GENL", "AL": "AUTO", "LE": "LAWE", "PO": "PUBO", "EP": "EMPL"}


def us(d):
    return d.strftime("%m/%d/%Y")


def money(v):
    return "%.2f" % v


def write_legacy(path, rec, mem):
    code = {m["no"]: m["code"] for m in mem}
    rows = [r for r in rec if r["system"] == "legacy"]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    heads = ["Claim No", "Member", "Coverage", "Date of Loss", "Date Tendered", "Date Closed", "Status",
             "Member Deductible", "Indemnity Paid", "Defense Paid", "Total Paid"]
    widths = [11, 8, 9, 12, 13, 12, 7, 17, 14, 13, 13]
    summary = []
    for ay in sorted({r["ay"] for r in rows}):
        ws = wb.create_sheet("AY%d" % ay)
        ws["A1"] = W.LEGACY_ADMIN
        ws["A2"] = "%s - Liability closed claims" % W.POOL
        ws["A3"] = "Accident year %d    Run date 07/02/2020" % ay
        ws["A1"].font = bold
        ws.append([])
        ws.append(heads)
        for c in ws[5]:
            c.font = bold
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        mine = sorted((r for r in rows if r["ay"] == ay), key=lambda r: r["legacy_no"])
        for r in mine:
            ws.append([r["legacy_no"], code[r["member"]], LEGACY_PART[r["part"]], us(r["date"]), us(r["tender"]),
                       us(r["close"]), "CL", r["ded"], r["paid_ind"], r["paid_def"],
                       round(r["paid_ind"] + r["paid_def"], 2)])
        for row in ws.iter_rows(min_row=6, max_row=ws.max_row):
            for c in row[7:]:
                c.number_format = "#,##0.00"
        for i, wd in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = wd
        ws.freeze_panes = "A6"
        summary.append((ay, len(mine), round(sum(r["paid_ind"] + r["paid_def"] for r in mine), 2)))
    ws = wb.create_sheet("Summary", 0)
    ws["A1"] = W.LEGACY_ADMIN
    ws["A2"] = "%s - Liability closed claims by accident year" % W.POOL
    ws["A3"] = "Claims closed through 06/30/2020    Run date 07/02/2020"
    ws["A1"].font = bold
    ws.append([])
    ws.append(["Accident Year", "Claims Closed", "Total Paid"])
    for c in ws[5]:
        c.font = bold
    for s in summary:
        ws.append(list(s))
    for row in ws.iter_rows(min_row=6, max_row=ws.max_row):
        row[2].number_format = "#,##0.00"
    for i, wd in enumerate((14, 14, 16), 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    wb.save(path)
    return len(rows)


def write_current(path, rec):
    rows = sorted((r for r in rec if r["system"] == "current" and r["status"] != "untendered"),
                  key=lambda r: r["claim_no"])
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["claim_no", "prior_claim_no", "member_no", "coverage", "loss_date", "tender_date", "status",
                    "closed_date", "deductible", "paid_indemnity", "paid_defense", "reserve_indemnity",
                    "reserve_defense"])
        for r in rows:
            w.writerow([r["claim_no"], r.get("legacy_no", "") if r["converted"] else "", r["member"],
                        CURRENT_PART[r["part"]], r["date"].isoformat(), r["tender"].isoformat(),
                        "Closed" if r["status"] == "closed" else "Open",
                        r["close"].isoformat() if r["status"] == "closed" else "", r["ded"], money(r["paid_ind"]),
                        money(r["paid_def"]), money(r["res_ind"]), money(r["res_def"])])
    return len(rows)


def write_log(path, log):
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["txn_id", "member_no", "unit_no", "txn_type", "effective_date", "entered_date", "vin",
                    "serial_no", "model_year", "make", "unit_class", "stated_value"])
        for r in log:
            w.writerow([r["txn"], r["member"], r["unit_no"], r["type"], r["effective"].isoformat(),
                        r["entered"].isoformat(), r["vin"], r["serial"], r["year"], r["make"], r["cls"],
                        "" if r["value"] is None else "%d" % r["value"]])
    return len(log)


def write_declarations(path, dec):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Declarations"
    heads = ["Member No", "Member Class", "Budget Basis", "Department", "FTE", "Payroll",
             "Total Operating Expenditures", "Submitted"]
    ws.append(heads)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    n = 0
    for d in dec:
        for i, ln in enumerate(d["lines"]):
            toe = d["toe"] if (not d["paper"] or i == 0) else None
            ws.append([d["member"], d["kind"], d["basis"], ln["department"], ln["fte"], ln["payroll"], toe,
                       dt.datetime.combine(d["submitted"], dt.time())])
            n += 1
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        row[4].number_format = "0.0"
        row[5].number_format = "#,##0"
        row[6].number_format = "#,##0"
        row[7].number_format = "mm/dd/yyyy"
    for i, wd in enumerate((10, 27, 12, 22, 8, 14, 16, 11), 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:H%d" % ws.max_row
    wb.save(path)
    return n


def founded(m):
    return m["founded"]


def write_roster(path, mem):
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["member_no", "member_code", "member_class", "joined", "withdrew"])
        for m in sorted(mem, key=lambda m: m["no"]):
            w.writerow([m["no"], m["code"], m["kind"], founded(m).isoformat(),
                        dt.date(m["left"], 12, 31).isoformat() if m["left"] is not None else ""])
    return len(mem)


def write_deductible_history(path, mem):
    n = 0
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["member_no", "effective_date", "deductible"])
        for m in sorted(mem, key=lambda m: m["no"]):
            for k, (eff, ded) in enumerate(m["deductibles"]):
                start = eff if k else max(eff, dt.date(W.FIT_YEARS[0], 1, 1))
                w.writerow([m["no"], start.isoformat(), ded])
                n += 1
    return n


def _noise(key):
    h = hashlib.sha256(key.encode()).digest()
    return 0.7 + (h[0] * 256 + h[1]) / 65535 * 0.75


def position(r, e):
    """The pool's paid and case reserve on an occurrence at an evaluation date."""
    if r["tender"] > e:
        return None
    if r["close"] <= e and r["status"] != "open":
        return r["pool_ult"], 0.0
    if e == W.AS_OF:
        return round(r["paid_ind"] + r["paid_def"], 2), round(r["res_ind"] + r["res_def"], 2)
    frac = ((e - r["tender"]).days / max(1, (r["close"] - r["tender"]).days)) ** 1.3
    paid = round(r["pool_ult"] * min(0.92, frac), 2)
    res = round(min(max(0.0, (r["pool_ult"] - paid) * _noise("%s%s%s" % (r["date"], r["member"], e))),
                    r["ret"] - r["ded"] - paid), 2)
    return paid, res


def write_triangles(path, rec):
    rows = [r for r in rec if r["status"] != "untendered"]
    n = 0
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["accident_year", "evaluation_date", "months", "tendered_count", "closed_count", "paid_to_date",
                    "incurred_to_date"])
        for ay in W.ALL_YEARS:
            evals = [dt.date(y, 12, 31) for y in range(ay, W.AS_OF.year)] + [W.AS_OF]
            for e in evals:
                t = c = 0
                paid = inc = 0.0
                for r in rows:
                    if r["ay"] != ay:
                        continue
                    pos = position(r, e)
                    if pos is None:
                        continue
                    t += 1
                    if r["close"] <= e and r["status"] != "open":
                        c += 1
                    paid += pos[0]
                    inc += pos[0] + pos[1]
                months = (e.year - ay) * 12 + e.month
                w.writerow([ay, e.isoformat(), months, t, c, money(paid), money(inc)])
                n += 1
    return n
