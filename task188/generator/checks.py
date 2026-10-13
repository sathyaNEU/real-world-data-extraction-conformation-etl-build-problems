"""Assertions on the built pack and goldens: the input gates; the three rules executed on the generator's clean truth
(R1 every pool ships with its size and result and no file states a rate, an index or a ranking; R2 each Layer 1 item
is neutral for the call and the graded bins; R3a the easy route, the minimum infection rate of every pool times the
CO2 catch, files Riverside South with every filed rule applied; R3b the decisive step alone files Westbrook and each
single correction stays wrong; R3c nothing planted); the ladder by name with margins and the golden's place on each;
every estimator variant in one bin; the semantic forks SF1 to SF13; the asks' golden against the truth and every stop
outside it for every zone; the separation of the asks from the call; the documents' silence; the realism sweep; and
the goldens against figures.json."""
import datetime as dt
import json
import math
import os
import re
import sqlite3
import zipfile

import openpyxl
from scipy.optimize import brentq

import analysis as A
import build as B
import cases as C
import deliver as DL
import docs as D
import pack as P
import records as R
import world as W

DAY = dt.timedelta(days=1)

N = [0]
HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)


def ok(cond, what):
    N[0] += 1
    if not cond:
        raise AssertionError(what)


def _pdf_text(path):
    from pypdf import PdfReader
    return re.sub(r"\s+", " ", " ".join(pg.extract_text() for pg in PdfReader(path).pages))


def jeffreys(pools):
    """The likelihood penalised by half the log of the information in p (a Bayesian mode, sometimes called Firth)."""
    if not any(x for m, x in pools):
        return 0.0

    def f(p):
        U, I, dI, k = A._parts(pools, p)
        return U + dI / (2.0 * I)
    return brentq(f, 1e-9, 0.5, xtol=1e-15)


EST = {"MLE": A.mle, "BC": A.bcmle, "JEF": jeffreys, "MIR": A.mir}


def lead(d, winner):
    rest = sorted((v for z, v in d.items() if z != winner), reverse=True)
    return d[winner] / rest[0] if rest[0] > 0 else 99.0


def pick(d):
    return max(d, key=d.get)


def place(d, z):
    return sorted(d, key=lambda k: -d[k]).index(z) + 1


def in_bin(vals, step, clear):
    b = {round(v / step) for v in vals}
    if len(b) != 1:
        return False
    c = b.pop() * step
    return min(min(v - (c - step / 2), (c + step / 2) - v) for v in vals) >= clear


class Truth:
    """The generator's own record of week 32: catches and pools as simulated, results as finally read."""

    def __init__(self, w):
        sim, flat = w["sim"], w["flat"]
        self.tn = [(r, c) for r, c in zip(sim["tn"], sim["catches"]) if r["year"] == 2026 and r["week"] == 32 and
                   r["origin"] == "field"]
        self.pools = [p for p in flat if p["year"] == 2026 and p["week"] == 32]

    def abundance(self, z, with_failed=False, with_males=False):
        tns = [(r, c) for r, c in self.tn if r["zone"] == z and r["trap_type"] == "CO2" and
               (r["status"] == "OK" or (with_failed and r["status"] in DL.FAILURES))]
        sexes = ("F", "M") if with_males else ("F",)
        return {sp: sum(n for r, c in tns for s, sex, n in c if s == sp and sex in sexes) / len(tns) for sp in W.VECTORS}

    def pl(self, z, which="CO2"):
        out = {sp: [] for sp in W.VECTORS}
        for p in self.pools:
            if p["zone"] == z and (which == "ALL" or p["trap_type"] == "CO2"):
                out[p["species"]].append((p["females"], p["final"] == "POS"))
        return out

    def index(self, z, est="MLE", which="CO2", genus=False):
        ab, pl = self.abundance(z), self.pl(z, which)
        f = EST[est]
        if genus:
            return sum(ab.values()) * f(pl["CXTA"] + pl["CXPI"])
        return sum(ab[sp] * f(pl[sp]) for sp in W.VECTORS)


def run(target, golden, w, meta_path, quiet=False):
    files = sorted(os.listdir(target))
    meta = json.load(open(meta_path))
    fig = json.load(open(os.path.join(HERE, "figures.json")))
    pk = DL.Pack(target)
    T = Truth(w)
    Z = W.ZONES
    prompt = open(os.path.join(TASK, "prompt.txt")).read()

    # ------------------------------------------------------------------ the input gates
    ok(len(files) >= 10, "ten or more files (%d)" % len(files))
    ok(len({f.rsplit(".", 1)[1] for f in files}) >= 3, "three or more formats")
    con = sqlite3.connect(os.path.join(target, R.DB))
    nrows = {t: con.execute("SELECT COUNT(*) FROM %s" % t).fetchone()[0] for t in ("catches", "pools", "trap_nights")}
    ok(max(nrows.values()) >= 25000, "a table of 25,000 rows or more (%s)" % nrows)
    ok(set(meta["distractor_files"]) == {P.WEATHER, P.BIRDS, P.FLIGHTS}, "three distractors declared")
    ok(not any(x in f.lower() for f in files for x in ("distractor", "golden", "answer", "trap", "decoy")),
       "no author noun in a file name")
    ok(files == sorted(B.FILES.values()), "the shipped set is the declared set")

    # ------------------------------------------------------------------ R1, the field's raw material ships
    cols = {t: [r[1] for r in con.execute("PRAGMA table_info(%s)" % t)] for t in
            ("sites", "trap_nights", "catches", "pools", "pool_tests", "species", "status_codes")}
    ok(cols["pools"] == ["pool_id", "tn_id", "species", "females", "lab", "sent"], "R1 every pool's size ships")
    ok(cols["pool_tests"] == ["pool_id", "test_no", "tested", "result", "ct"], "R1 every test's result ships")
    allcols = " ".join(" ".join(v) for v in cols.values()).lower()
    for wd in ("rate", "mir", "mle", "index", "infected", "positive_share", "trap_type_of_pool"):
        ok(re.search(r"\b%s\b" % wd, allcols) is None, "R1 no derived column (%s)" % wd)
    ok(len(pk.county) == len(B.county_week32(w["flat"])), "R1 the county report lists every week-32 pool (%d)" %
       len(pk.county))
    texts = {}
    for f in files:
        if f.endswith((".csv", ".txt", ".geojson")):
            texts[f] = open(os.path.join(target, f), encoding="utf-8").read()
        elif f.endswith(".pdf"):
            texts[f] = _pdf_text(os.path.join(target, f))
    for f, t in texts.items():
        lo = t.lower()
        for wd in ("infection rate", "minimum infection", "maximum likelihood", "vector index", "per 1,000",
                   "likelihood", "bias", "infected per"):
            ok(wd not in lo, "R1 %s states no rate (%s)" % (f, wd))
    # R3c nothing planted: no file names a zone for this week's flight or ranks the zones on the question
    ok("2026-08-20" not in texts[P.FLIGHTS], "R3c the flight log stops before this week's flight")
    for f in (D.PROTOCOL, D.DICTIONARY, D.COUNTY, D.CONTRACT):
        lo = texts[f].lower()
        for wd in ("westbrook should", "recommend", "highest risk", "hot spot", "priority zone"):
            ok(wd not in lo, "R3c %s carries no recommendation (%s)" % (f, wd))
    cty = texts[D.COUNTY]
    ok(not any(z in cty for z in Z), "R3c the county report names no zone")
    pro = texts[D.PROTOCOL].lower()
    for wd in ("likelihood", "minimum infection", "pool size", "estimat", "infection rate", "index", "gravid females "
               "carry", "older"):
        ok(wd not in pro, "R3 the protocol is silent on %s" % wd)
    plo = prompt.lower()
    for wd in ("likelihood", "gravid", "minimum", "pool size", "rate of infection", "mir", "mle", "host-seeking",
               "blood meal"):
        ok(re.search(r"\b%s" % wd, plo) is None, "R3 the prompt is silent on %s" % wd)

    # ------------------------------------------------------------------ R3, the analyst test on the clean truth
    cell = {}
    for key, est, which in (("default", "MIR", "ALL"), ("co2_mir", "MIR", "CO2"), ("all_mle", "MLE", "ALL"),
                            ("co2_mle", "MLE", "CO2")):
        cell[key] = {z: T.index(z, est, which) for z in Z}
    ok(pick(cell["default"]) == "Riverside South", "R3a the easy route on the clean truth files Riverside South (%s)"
       % pick(cell["default"]))
    ok(pick(cell["co2_mle"]) == "Westbrook", "R3b the decisive step alone files Westbrook")
    ok(pick(cell["co2_mir"]) == "South Prairie", "R3b the CO2 pools alone file South Prairie")
    ok(pick(cell["all_mle"]) == "Old Town", "R3b the likelihood estimate alone files Old Town")
    for key in cell:
        b = {z: pk.index(z, *{"default": ("MIR", "ALL"), "co2_mir": ("MIR", "CO2"), "all_mle": ("MLE", "ALL"),
                               "co2_mle": ("MLE", "CO2")}[key]) for z in Z}
        ok(all(abs(b[z] - cell[key][z]) < 1e-9 for z in Z), "R2 the pack's %s cells equal the truth's (C1)" % key)

    # ------------------------------------------------------------------ the ladder, built values
    margins = {"default": (1.25, "Riverside South"), "co2_mir": (1.15, "South Prairie"), "all_mle": (1.15, "Old Town"),
               "co2_mle": (1.25, "Westbrook")}
    for key, (m, z) in margins.items():
        ok(pick(cell[key]) == z and lead(cell[key], z) >= m, "ladder %s files %s by %.2fx or more (%.3f)" % (
            key, z, m, lead(cell[key], z)))
    ok(place(cell["default"], "Westbrook") == 4, "Westbrook fourth on the default")
    ok(place(cell["co2_mir"], "Westbrook") == 2 and place(cell["all_mle"], "Westbrook") == 3, "Westbrook second and "
                                                                                              "third on the rungs")
    dom = (cell["co2_mle"]["Westbrook"] - cell["co2_mle"]["Riverside South"]) / (
        cell["default"]["Riverside South"] - cell["default"]["Westbrook"])
    ok(dom >= 1.2, "dominance: Westbrook's lead under the decisive method over Riverside South's carried advantage "
                   "(%.2f)" % dom)
    # every reflex lands on a named wrong zone
    ab = {z: sum(T.abundance(z).values()) for z in Z}
    ok(pick(ab) == "North Bottoms", "reflex: catch alone files North Bottoms")
    rate = {z: A.mle(T.pl(z)["CXTA"] + T.pl(z)["CXPI"]) for z in Z}
    ok(pick(rate) == "Maple Ridge" and lead(rate, "Maple Ridge") >= 1.15, "reflex: infection alone files Maple Ridge "
                                                                          "(%.2f)" % lead(rate, "Maple Ridge"))
    npos = {z: sum(x for sp in W.VECTORS for m, x in T.pl(z, "ALL")[sp]) for z in Z}
    ok(pick(npos) == "Old Town", "reflex: positive pools counted files Old Town")
    share = {z: npos[z] / sum(len(T.pl(z, "ALL")[sp]) for sp in W.VECTORS) for z in Z}
    ok(pick(share) == "Old Town", "reflex: share of pools positive files Old Town")
    ok(pick(fig["cases_by_zone"]) == "Coulee Hills", "reflex: cases files Coulee Hills")
    gm = {z: T.index(z, "MIR", "ALL", genus=True) for z in Z}
    ok(pick(gm) == "Riverside South", "reflex: one Culex minimum rate over every pool files Riverside South")
    ok(pick({z: T.index(z, "MIR", "CO2", genus=True) for z in Z}) == "South Prairie", "reflex: one Culex minimum rate "
                                                                                      "over the CO2 pools files South "
                                                                                      "Prairie")
    ok(pick({z: T.index(z, "MLE", "ALL", genus=True) for z in Z}) == "Old Town", "reflex: one Culex likelihood rate "
                                                                                 "over every pool files Old Town")
    # gravid-trap catches added to the abundance (read past the prompt's CO2 trap-night)
    def ab_all(z):
        tns = [(r, c) for r, c in T.tn if r["zone"] == z and r["status"] == "OK"]
        return {sp: sum(n for r, c in tns for s, sex, n in c if s == sp and sex == "F") / len(tns) for sp in W.VECTORS}
    v = {z: sum(ab_all(z)[sp] * A.mle(T.pl(z, "ALL")[sp]) for sp in W.VECTORS) for z in Z}
    ok(pick(v) == "Old Town", "a trap-night of either kind with every pool files Old Town (%s)" % pick(v))

    # ------------------------------------------------------------------ one bin under every estimator
    call, second = fig["call"], fig["runner_up"]
    cw = sum(T.abundance(call).values())
    vals_c, vals_s, flips = [], [], []
    for est in ("MLE", "BC", "JEF"):
        for genus in (False, True):
            d = {z: T.index(z, est, "CO2", genus) for z in Z}
            ok(pick(d) == call and lead(d, call) >= 1.25, "%s%s files %s by 1.25x or more (%.3f)" % (
                est, " genus" if genus else "", call, lead(d, call)))
            ok(sorted(d, key=lambda k: -d[k])[1] == second, "%s%s: the runner-up is %s" % (est, genus, second))
            vals_c.append(d[call])
            vals_s.append(d[second])
            flips.append(d[second] / cw * 1000.0)
    ok(in_bin(vals_c, 0.1, 0.010), "SF5 SF6 %s's figure in one bin under every estimator (%s)" % (
        call, [round(x, 3) for x in vals_c]))
    ok(in_bin(vals_s, 0.1, 0.012), "SF5 SF6 %s's figure in one bin (%s)" % (second, [round(x, 3) for x in vals_s]))
    ok(in_bin(flips, 1.0, 0.15), "SF5 SF6 the flip rate in one whole number (%s)" % [round(x, 2) for x in flips])
    ok(fig["call_figure"] == round(vals_c[0], 1) and fig["runner_up_figure"] == round(vals_s[0], 1) and
       fig["flip_rate_per_1000"] == int(round(flips[0])), "figures.json carries the bins")
    ok(fig["bias_corrected"]["call"] == fig["call_figure"] and fig["bias_corrected"]["runner_up"] ==
       fig["runner_up_figure"] and fig["bias_corrected"]["flip"] == fig["flip_rate_per_1000"], "SF5 the bias-corrected "
                                                                                               "figures agree")
    # the flip point is real: at that infection the call's figure equals the runner-up's
    ok(abs(fig["flip_rate_exact"] / 1000.0 * cw - cell["co2_mle"][second]) < 1e-4, "SF10 the flip rate reproduces")

    ok(all(any(x for sp in W.VECTORS for m, x in T.pl(z, "ALL")[sp]) for z in Z), "the prompt's premise: every zone "
                                                                                    "returned a positive pool in week 32")

    # ------------------------------------------------------------------ the semantic forks
    pick_up = {z: [r for r, c in T.tn if r["zone"] == z] for z in Z}
    ok(all(DL.WEEK32[0] <= r["pickup_date"].isoformat() <= DL.WEEK32[1] for z in Z for r in pick_up[z]),
       "SF1 every week-32 trap-night is picked up inside the week (C1)")
    ok(all(dt.date(2026, 8, 9) <= r["set_date"] <= dt.date(2026, 8, 15) for z in Z for r in pick_up[z]), "SF1 and set")
    ok(sum(n for r, c in T.tn for s, sex, n in c if s == "CXSP") == 0, "SF2 no unidentified Culex in week 32 (C1)")
    d3 = {}
    for z in Z:
        ab3 = T.abundance(z, with_failed=True)
        d3[z] = sum(ab3[sp] * A.mle(T.pl(z)[sp]) for sp in W.VECTORS)
    ok(pick(d3) == call and round(d3[call], 1) == fig["call_figure"] and round(d3[second], 1) ==
       fig["runner_up_figure"], "SF3 failed traps counted: same call and figures (C1 on the two zones)")
    ok(all(r["status"] == "OK" for r, c in T.tn if r["zone"] in (call, second) and r["trap_type"] == "CO2"),
       "SF3 SF7 no failed trap in %s or %s in week 32" % (call, second))
    inc = [p for p in T.pools if p["result"] == "INC"]
    ok(not any(p["zone"] in (call, second) for p in inc), "SF4 no inconclusive pool in %s or %s (C1)" % (call, second))
    for as_pos in (True, False):
        d4 = {}
        for z in Z:
            abz = T.abundance(z)
            pl = {sp: [(p["females"], (p["result"] == "POS") or (as_pos and p["result"] == "INC")) for p in T.pools
                       if p["zone"] == z and p["trap_type"] == "CO2" and p["species"] == sp] for sp in W.VECTORS}
            d4[z] = sum(abz[sp] * A.mle(pl[sp]) for sp in W.VECTORS)
        ok(pick(d4) == call, "SF4 first results as final (inconclusive read %s): same call" % (
            "positive" if as_pos else "negative"))
    ok(True, "SF5 and SF6 asserted under the bins")
    for z in (call, second):
        tns = [(r, c) for r, c in T.tn if r["zone"] == z and r["trap_type"] == "CO2" and r["status"] == "OK"]
        persite = {}
        for r, c in tns:
            persite.setdefault(r["site_id"], []).append(sum(n for s, sex, n in c if s in W.VECTORS and sex == "F"))
        ok(all(len(v_) == 2 for v_ in persite.values()), "SF7 every %s site ran both nights, so the mean of site means "
                                                         "is the total over trap-nights (C1)" % z)
    wb = openpyxl.load_workbook(os.path.join(target, P.REGISTER), read_only=True)
    reg = {r[0]: r[1] for r in wb["Sites"].iter_rows(min_row=2, values_only=True)}
    ok(all(reg[s] == pk.sites[s]["zone"] for s in pk.sites), "SF8 the register and the database agree on every "
                                                              "site's zone")
    ok(all(pk.zone_of_section[pk.sites[s]["section"]] == pk.sites[s]["zone"] for s in pk.sites), "SF8 every site sits "
                                                                                                "in its zone's sections")
    ok("its figure" in prompt and "runner-up" in prompt, "SF9 the runner-up's figure is the same figure")
    ok("per 1,000 Culex females" in prompt, "SF10 one rate across the genus")
    males = {z: sum(n for r, c in T.tn if r["zone"] == z and r["trap_type"] == "CO2" and r["status"] == "OK"
                    for s, sex, n in c if sex == "M") for z in Z}
    ok(males[call] > 0, "SF11 the CO2 traps catch males")
    ok(all(p["females"] > 0 for p in T.pools), "SF11 pools hold females")
    d11 = {}
    for z in Z:
        ab11 = T.abundance(z, with_males=True)
        d11[z] = sum(ab11[sp] * A.mle(T.pl(z)[sp]) for sp in W.VECTORS)
    ok(pick(d11) == call, "SF11 males added: same call (ruled out by the prompt's females)")
    ok("Culex females" in prompt, "SF11 the prompt counts females")
    ok("one block" in texts[D.CONTRACT].lower() and "not split" in texts[D.CONTRACT].lower(), "SF12 one block, one zone")
    ok(True, "SF13 asserted under the asks")

    # ------------------------------------------------------------------ the asks
    rows = w["cases"]
    tc = C.truth(rows)
    for z in Z:
        ok(fig["cases_by_zone"][z] == len(tc[z]), "A golden equals the truth for %s" % z)
    gold_a = fig["cases_by_zone"]
    stops_a = {}
    lab, lab_all, ev, sec_all = {z: 0 for z in Z}, {z: 0 for z in Z}, {z: 0 for z in Z}, {z: 0 for z in Z}
    seen, seen2 = set(), set()
    for r in rows:
        lab_all[r["label"]] += 1
        sec_all[W.zone2026(*r["_xy"])] += 1
        excluded = r["pres"] == "Asymptomatic blood donor" or r["exposure"] != "Rindal County" or r["cls"] == "Not a case"
        if not excluded and r["onset"] and W.SEASON_2026[0] <= r["onset"] <= W.SEASON_2026[1] and r["state_id"] not in seen:
            seen.add(r["state_id"])
            lab[r["label"]] += 1
        if not excluded and W.SEASON_2026[0] <= r["event"] <= W.SEASON_2026[1] and r["state_id"] not in seen2:
            seen2.add(r["state_id"])
            ev[W.zone2026(*r["_xy"])] += 1
    stops_a = {"label": lab, "label, every row": lab_all, "map, every row": sec_all}
    for k, d in stops_a.items():
        for z in Z:
            ok(d[z] != gold_a[z], "A stop '%s' misses %s (%d against %d)" % (k, z, d[z], gold_a[z]))
    ok({z for z in Z if ev[z] != gold_a[z]} == {"Coulee Hills", "South Prairie"}, "A event date for onset misses the "
                                                                                  "two zones with May onsets")
    for kind, zs in (("donor", {"North Bottoms", "Westbrook", "South Prairie"}), ("travel", {"Old Town", "Maple Ridge"}),
                     ("ruled_out", {"Maple Ridge", "Riverside South", "Airport East"}),
                     ("duplicate", {"Coulee Hills", "Westbrook"}), ("may", {"Coulee Hills", "South Prairie"})):
        ok({W.zone2026(*r["_xy"]) for r in rows if r["_kind"] == kind} == zs, "A hazard %s sits in %s" % (kind, zs))
    moved = [r for r in rows if r["label"] != W.zone2026(*r["_xy"])]
    ok(len(moved) == 8 and all(r["_kind"] == "case" for r in moved), "A eight cases carry a 2025 zone label")
    ok("March 2025" in str(openpyxl.load_workbook(os.path.join(target, C.LOG), read_only=True)["Read me"]["A9"].value),
       "A the log's read-me dates its zone list")
    # the enhanced nights follow Maple Ridge's first positive pool of the season, as protocol 3.5 has it
    mr_first = min(p["collected"] for p in w["flat"] if p["zone"] == "Maple Ridge" and p["year"] == 2026 and
                   p["final"] == "POS")
    ok(mr_first + dt.timedelta(days=3) < min(W.ENHANCED) and min(W.ENHANCED) - mr_first <= dt.timedelta(days=14) and
       max(W.ENHANCED) - min(W.ENHANCED) < dt.timedelta(days=14), "protocol 3.5: Maple Ridge's extra nights follow its "
                                                                  "first positive pool, within two weeks")
    first_pos = {}
    for p in w["flat"]:
        if p["year"] == 2026 and p["final"] == "POS":
            first_pos[p["zone"]] = min(first_pos.get(p["zone"], p["collected"]), p["collected"])
    ok(min(first_pos, key=first_pos.get) == "Maple Ridge", "Maple Ridge had the season's first positive pool")
    ok("redrawn on 1 May 2026" in texts[D.PROTOCOL], "A the protocol dates the redraw")
    ok("Issued 27 April 2026, revised 26 June 2026" in texts[D.PROTOCOL], "the protocol's revision postdates the "
                                                                          "migration and the redraw it describes")
    # B
    gold_r, gold_l = fig["co2_trap_nights_run"], fig["co2_trap_nights_lost"]
    sim = w["sim"]
    truth_b = {z: [0, 0] for z in Z}
    for r in sim["tn"]:
        if r["origin"] == "reimport" or r["trap_type"] != "CO2" or not (W.SEASON_2026[0] <= r["set_date"] <=
                                                                       W.SEASON_2026[1]):
            continue
        if r["status"] == "OK":
            truth_b[r["zone"]][0] += 1
        elif r["status"] in DL.FAILURES:
            truth_b[r["zone"]][1] += 1
    for z in Z:
        ok([gold_r[z], gold_l[z]] == truth_b[z], "B golden equals the truth for %s" % z)
    ok(fig["reimport_copies"] == sum(1 for r in sim["tn"] if r["origin"] == "reimport") == 206, "B the copies found "
                                                                                                 "are the second load's")
    # every copy is ruled out by one shipped sentence: it sits on a night the protocol schedules no trap (3.3 Monday
    # and Tuesday, 3.5 the enhanced nights), and no copy lands on a night a site actually ran, so no record pair asks
    # which of two readings of one night to keep
    nights = {(r["site_id"], r["set_date"]) for r in sim["tn"] if r["origin"] != "reimport"}
    for r in sim["tn"]:
        if r["origin"] != "reimport":
            continue
        ok((r["site_id"], r["set_date"]) not in nights, "B copy %s on %s lands on no night the site ran" % (
            r["site_id"], r["set_date"]))
        ok(r["set_date"].weekday() not in (0, 1) and r["set_date"] not in W.ENHANCED,
           "B copy %s on %s falls on a day the protocol schedules no trap" % (r["site_id"], r["set_date"]))
        same = [(q, c) for q, c in zip(sim["tn"], sim["catches"]) if q["site_id"] == r["site_id"] and
                q["set_date"] == r["set_date"] - DAY and q["origin"] == "field"]
        ok(len(same) == 1, "B copy %s repeats one night of its site" % r["site_id"])
    ok({pk.tn[t]["entered"][:10] for t in pk.copies()} == {"2026-06-22"}, "B every copy was entered on 22 June")
    ok("22 June 2026" in texts[D.DICTIONARY], "B the dictionary dates the import run")
    ok("A site runs one trap a night" in texts[D.PROTOCOL], "SF14 the protocol pins one trap a site a night")
    ok("loaded again on 22 June 2026" in texts[D.PROTOCOL], "SF14 the protocol records the second load")
    every = pk.season_trap_nights(drop=set())
    seen_k, dup = set(), set()
    for t, r in sorted(pk.tn.items()):
        k = (r["site"], r["set"], r["status"], tuple(sorted(pk.catches.get(t, []))))
        if k in seen_k:
            dup.add(t)
        seen_k.add(k)
    exact = pk.season_trap_nights(drop=dup)
    for z in Z:
        ok(every[z][0] != gold_r[z] and every[z][1] != gold_l[z], "B stop 'every row' misses %s on both" % z)
        ok(exact[z][0] != gold_r[z] and exact[z][1] != gold_l[z], "B stop 'exact duplicates' misses %s on both" % z)
    mt = {t for t, r in pk.tn.items() if dt.date.fromisoformat(r["set"]).weekday() not in (0, 1)}
    over = pk.season_trap_nights(drop=mt)
    ok({z for z in Z if over[z][0] != gold_r[z]} == {"Maple Ridge"}, "B dropping every night off Monday and Tuesday "
                                                                     "clears the copies and over-cleans Maple Ridge")
    nsl = {z: [0, 0] for z in Z}
    cp = pk.copies()
    for t, r in pk.tn.items():
        if t in cp or not (DL.SEASON[0] <= r["set"] <= DL.SEASON[1]) or pk.sites[r["site"]]["trap_type"] != "CO2":
            continue
        nsl[pk.sites[r["site"]]["zone"]][0 if r["status"] == "OK" else 1] += 1
    ok(all(nsl[z][1] != gold_l[z] for z in Z), "B stop 'not set counted as lost' misses every zone")
    # the chart's bars
    for z in Z:
        ok(sum(fig["cases_by_week"][z]) == gold_a[z], "chart bars sum to %s's cases" % z)
    # separation: nothing an ask turns on sits in the main population
    t32 = [r for r in sim["tn"] if r["year"] == 2026 and r["week"] == 32]
    ok(not any(r["origin"] != "field" for r in t32), "separation: no copy or enhanced night in week 32")
    ok(not any(r["status"] == "NS" for r in t32), "separation: no trap left unset in week 32")
    ok(all(r["set_date"] > W.MIGRATION for r in t32), "separation: week 32 is after the migration")

    # ------------------------------------------------------------------ the documents and the dates
    late = []
    for f in files:
        if f.endswith((".csv", ".txt", ".geojson")):
            if any(x > W.AS_OF.isoformat() for x in set(re.findall(r"\b20\d\d-\d\d-\d\d\b", texts[f]))):
                late.append(f)
    ok(late == [], "no record runs past the extract (%s)" % late)
    mx = con.execute("SELECT MAX(set_date), MAX(pickup_date), MAX(entered) FROM trap_nights").fetchone()
    ok(mx[0] <= "2026-08-11" and mx[1] <= "2026-08-12" and mx[2] < "2026-08-17 08:20", "database dates before the "
                                                                                        "extract (%s)" % (mx,))
    ok(con.execute("SELECT MAX(tested) FROM pool_tests").fetchone()[0] <= "2026-08-16", "tests before the extract")
    man = texts[P.MANIFEST]
    ok({f for f in files if f in man} == set(files) - {P.MANIFEST}, "the manifest names every other file")
    dic = texts[D.DICTIONARY]
    ok(all(t in dic for t in cols), "the dictionary names every table")

    # ------------------------------------------------------------------ the realism sweep
    ok(len(set(gold_r.values())) >= 7, "trap-nights run vary by zone")
    ok(len({len(c) for r, c in T.tn}) > 3, "catch rows vary by trap-night")
    sizes = [p["females"] for p in T.pools if p["lab"] == "district"]
    ok(sizes.count(50) > len(sizes) / 3 and len(set(sizes)) > 30, "district pools: many full, the rest of every size")
    allt = " ".join(texts.values())
    for k in ("surv", "ops", "pilot"):
        ok(W.PEOPLE[k][0] in allt, "persona %s in the pack" % W.PEOPLE[k][0])
    for f in files:
        if f.endswith(".xlsx"):
            with zipfile.ZipFile(os.path.join(target, f)) as z:
                core = z.read("docProps/core.xml").decode()
                app = z.read("docProps/app.xml").decode()
            ok("openpyxl" not in core and "Microsoft Excel" in app, "container scrubbed: %s" % f)
        if f.endswith(".pdf"):
            b = open(os.path.join(target, f), "rb").read()
            ok(b"ReportLab" not in b and b"reportlab" not in b, "container scrubbed: %s" % f)
    for f, t in texts.items():
        ok("—" not in t and "–" not in t, "no dash in %s" % f)

    # ------------------------------------------------------------------ the goldens agree with figures.json
    ok(sorted(os.listdir(golden)) == sorted([DL.MEMO, DL.WORKBOOK, DL.CHART]), "the goldens are the three files asked")
    mt_ = _pdf_text(os.path.join(golden, DL.MEMO))
    ok("Aerial block for Thursday 20 August 2026: %s" % call in mt_, "the memo's title is the call")
    ok(("an estimated %.1f" % fig["call_figure"]) in mt_ and ("%s at %.1f" % (second, fig["runner_up_figure"])) in mt_,
       "the memo states both figures")
    ok(("below %d per 1,000" % fig["flip_rate_per_1000"]) in mt_, "the memo states the flip rate")
    ok("—" not in mt_ and "–" not in mt_, "no dash in the memo")
    gw = openpyxl.load_workbook(os.path.join(golden, DL.WORKBOOK))
    ok(gw.sheetnames == ["Zones", "Cases by week", "Notes"], "the workbook's sheets")
    for r, z in zip(gw["Zones"].iter_rows(min_row=2, max_row=9, values_only=True), Z):
        ok(list(r) == [z, gold_a[z], gold_r[z], gold_l[z]], "the workbook's row for %s" % z)
    for r in gw["Notes"].iter_rows(values_only=True):
        ok(r[0] is None or ("—" not in r[0] and "–" not in r[0]), "no dash in the workbook notes")
    png = open(os.path.join(golden, DL.CHART), "rb").read()
    ok(png[:8] == b"\x89PNG\r\n\x1a\n" and b"matplotlib" not in png.lower(), "the chart is a clean PNG")
    con.close()
    if not quiet:
        print("checks: %d assertions passed (call %s %.4f, runner-up %s %.4f, flip %.2f per 1,000; default %s by %.3f, "
              "Westbrook fourth; CO2 minimum rate %s by %.3f; likelihood on every pool %s by %.3f; decisive %s by %.3f; "
              "dominance %.2f)" % (N[0], call, cell["co2_mle"][call], second, cell["co2_mle"][second],
                                   fig["flip_rate_exact"], pick(cell["default"]), lead(cell["default"], "Riverside South"),
                                   pick(cell["co2_mir"]), lead(cell["co2_mir"], "South Prairie"), pick(cell["all_mle"]),
                                   lead(cell["all_mle"], "Old Town"), call, lead(cell["co2_mle"], call), dom))
