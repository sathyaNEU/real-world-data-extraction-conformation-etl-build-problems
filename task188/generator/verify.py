"""Independent verifier: reads only the shipped pack and the goldens, recomputes every graded figure by its own route
(SQL over the field database, pypdf over the county report, the pooled likelihood maximised directly rather than
solved from its score) and compares them with the goldens.

    python3 task188/generator/verify.py [target] [golden]"""
import datetime as dt
import math
import os
import re
import sqlite3
import sys

from scipy.optimize import minimize_scalar

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
ZONES = ["North Bottoms", "Maple Ridge", "Westbrook", "Old Town", "Riverside South", "South Prairie", "Coulee Hills",
         "Airport East"]


def loglik(p, pools):
    s = 0.0
    for m, pos in pools:
        q = (1.0 - p) ** m
        s += math.log(1.0 - q) if pos else math.log(q)
    return s


def mle(pools):
    if not any(pos for m, pos in pools):
        return 0.0
    r = minimize_scalar(lambda p: -loglik(p, pools), bounds=(1e-8, 0.3), method="bounded",
                        options={"xatol": 1e-13, "maxiter": 2000})
    return r.x


def county(path):
    """pypdf gives the report's table one cell a line: walk the cells after each accession number."""
    from pypdf import PdfReader
    cells = [c.strip() for pg in PdfReader(path).pages for c in pg.extract_text().splitlines()]
    rows, retest = {}, {}
    table2 = False
    i = 0
    while i < len(cells):
        c = cells[i]
        if c.startswith("Table 2"):
            table2 = True
        if re.fullmatch(r"RCPHL-26-\d{4}", c):
            if not table2:
                site, coll, sp, nn, ct, res = cells[i + 1:i + 7]
                mm, dd, yy = coll.split("/")
                rows[c] = dict(site=site, collected="%s-%s-%s" % (yy, mm, dd),
                               species={"Culex tarsalis": "CXTA", "Culex pipiens": "CXPI"}[sp], n=int(nn), result=res)
                i += 7
                continue
            retest[c] = cells[i + 3]
            i += 4
            continue
        i += 1
    for k, r in rows.items():
        r["positive"] = (retest[k] if r["result"] == "Inconclusive" else r["result"]) == "Positive"
    return rows


def main(target=None, golden=None):
    target = target or os.path.join(TASK, "target")
    golden = golden or os.path.join(TASK, "golden")
    db = [f for f in os.listdir(target) if f.endswith(".sqlite")][0]
    con = sqlite3.connect(os.path.join(target, db))
    pdf = [f for f in os.listdir(target) if f.startswith("RCPHL") and f.endswith(".pdf")][0]
    cty = county(os.path.join(target, pdf))
    n = [0]

    def ok(c, what):
        n[0] += 1
        if not c:
            raise AssertionError("verify: " + what)

    # week 32: females per functioning CO2 trap-night, by zone and species
    ab = {}
    for z, sp, f, tn in con.execute("""
            WITH t AS (SELECT n.tn_id, s.zone FROM trap_nights n JOIN sites s USING (site_id)
                       WHERE n.set_date BETWEEN '2026-08-09' AND '2026-08-15' AND s.trap_type = 'CO2' AND n.status = 'OK')
            SELECT t.zone, sp.code, COALESCE((SELECT SUM(c.count) FROM catches c JOIN t t2 USING (tn_id)
                                              WHERE t2.zone = t.zone AND c.species = sp.code AND c.sex = 'F'), 0),
                   COUNT(DISTINCT t.tn_id)
            FROM t, species sp WHERE sp.code IN ('CXTA', 'CXPI') GROUP BY t.zone, sp.code"""):
        ab[(z, sp)] = f / tn
    # week-32 CO2 pools with their final results: district results in the database, county results in the report
    final = {}
    for pid, res in con.execute("""SELECT pool_id, result FROM pool_tests t WHERE test_no =
                                   (SELECT MAX(test_no) FROM pool_tests u WHERE u.pool_id = t.pool_id)"""):
        final[pid] = res == "POS"
    pools = {}
    for pid, z, sp, m, typ in con.execute("""
            SELECT p.pool_id, s.zone, p.species, p.females, s.trap_type FROM pools p JOIN trap_nights n USING (tn_id)
            JOIN sites s USING (site_id) WHERE n.set_date BETWEEN '2026-08-09' AND '2026-08-15'"""):
        pos = final[pid] if pid in final else cty[pid]["positive"]
        if pid in cty:
            ok(cty[pid]["n"] == m and cty[pid]["species"] == sp, "county report and database agree on %s" % pid)
        if typ == "CO2":
            pools.setdefault((z, sp), []).append((m, pos))
    ok(len([p for p in cty]) == con.execute("""SELECT COUNT(*) FROM pools p JOIN trap_nights n USING (tn_id)
            WHERE p.lab = 'RCPHL' AND n.set_date BETWEEN '2026-08-09' AND '2026-08-15'""").fetchone()[0],
       "the county report holds every week-32 county pool")
    idx = {z: sum(ab[(z, sp)] * mle(pools.get((z, sp), [])) for sp in ("CXTA", "CXPI")) for z in ZONES}
    order = sorted(ZONES, key=lambda z: -idx[z])
    call, second = order[0], order[1]
    flip = idx[second] / (ab[(call, "CXTA")] + ab[(call, "CXPI")]) * 1000.0
    ok(idx[call] / idx[second] > 1.2, "the call leads by more than 1.2x")

    # ask A: locally acquired symptomatic cases by onset, 1 June to 15 August, one per state case ID, by 2026 zone
    from openpyxl import load_workbook
    reg = [f for f in os.listdir(target) if "register" in f][0]
    zs = {r[0]: r[4] for r in load_workbook(os.path.join(target, reg), read_only=True)["Zone sections"].iter_rows(
        min_row=2, values_only=True)}
    log = [f for f in os.listdir(target) if "case_log" in f][0]
    rows = list(load_workbook(os.path.join(target, log), read_only=True)["Cases"].iter_rows(values_only=True))
    h = {k: i for i, k in enumerate(rows[0])}
    cases, ids = {z: 0 for z in ZONES}, set()
    weeks = {z: {} for z in ZONES}
    for r in rows[1:]:
        if r[h["Classification"]] not in ("Confirmed", "Probable") or r[h["Exposure"]] != "Rindal County":
            continue
        on = r[h["Onset date"]]
        if on is None or not (dt.datetime(2026, 6, 1) <= on <= dt.datetime(2026, 8, 15)) or r[h["State case ID"]] in ids:
            continue
        ids.add(r[h["State case ID"]])
        z = zs[r[h["Residence section"]]]
        cases[z] += 1
        jan4 = dt.datetime(2026, 1, 4)
        wk = (on - jan4).days // 7 + 1
        weeks[z][wk] = weeks[z].get(wk, 0) + 1

    # ask B: CO2 trap-nights run and lost, 1 June to 15 August, the re-imported second copies counted once
    fp = {}
    for t, cat in con.execute("""SELECT tn_id, GROUP_CONCAT(species || sex || count, ',') FROM
                                 (SELECT * FROM catches ORDER BY tn_id, species, sex) GROUP BY tn_id"""):
        fp[t] = cat
    recs = list(con.execute("""SELECT n.tn_id, n.site_id, n.set_date, n.status, n.entered, s.zone, s.trap_type
                               FROM trap_nights n JOIN sites s USING (site_id)"""))
    key = {(r[1], r[2], r[3], fp.get(r[0], "")): r[0] for r in recs if not r[4].startswith("2026-06-22")}
    copies = set()
    for r in recs:
        if r[4].startswith("2026-06-22"):
            prev = (dt.date.fromisoformat(r[2]) - dt.timedelta(days=1)).isoformat()
            if (r[1], prev, r[3], fp.get(r[0], "")) in key:
                copies.add(r[0])
    ok(len(copies) == sum(1 for r in recs if r[4].startswith("2026-06-22")), "every re-imported record has its "
                                                                              "original a day earlier")
    own = {(r[1], r[2]) for r in recs if r[0] not in copies}
    for t, site, d, st, ent, z, typ in recs:
        if t not in copies:
            continue
        wd = dt.date.fromisoformat(d).weekday()
        enh = z == "Maple Ridge" and d in ("2026-06-17", "2026-06-18", "2026-06-24", "2026-06-25")
        ok((site, d) not in own and wd not in (0, 1) and not enh, "copy %s on %s falls on a day no trap is scheduled "
                                                                  "and on no night the site ran" % (site, d))
    run, lost = {z: 0 for z in ZONES}, {z: 0 for z in ZONES}
    for t, site, d, st, ent, z, typ in recs:
        if t in copies or typ != "CO2" or not ("2026-06-01" <= d <= "2026-08-15"):
            continue
        if st == "OK":
            run[z] += 1
        elif st in ("FAN", "CO2", "BAT", "STOLEN"):
            lost[z] += 1

    # the goldens
    wb = load_workbook(os.path.join(golden, "zone_season_2026.xlsx"), read_only=True)
    got = {r[0]: r[1:4] for r in wb["Zones"].iter_rows(min_row=2, values_only=True)}
    for z in ZONES:
        ok(tuple(got[z]) == (cases[z], run[z], lost[z]), "workbook row %s %s against %s" % (z, got[z], (cases[z], run[z],
                                                                                                         lost[z])))
    wk = {r[0]: r[1:] for r in wb["Cases by week"].iter_rows(min_row=2, values_only=True)}
    for z in ZONES:
        ok(list(wk[z][:-1]) == [weeks[z].get(w, 0) for w in range(22, 33)], "cases by week for %s" % z)
    from pypdf import PdfReader
    memo = re.sub(r"\s+", " ", " ".join(pg.extract_text() for pg in PdfReader(os.path.join(golden,
                                                                                          "aerial_spray_recommendation.pdf")).pages))
    ok("Fly %s." % call in memo, "the memo names %s" % call)
    ok("an estimated %.1f" % idx[call] in memo, "the memo's figure %.1f" % idx[call])
    ok("%s at %.1f" % (second, idx[second]) in memo, "the memo's runner-up %s %.1f" % (second, idx[second]))
    ok("below %d per 1,000" % round(flip) in memo, "the memo's flip rate %d" % round(flip))
    ok(os.path.getsize(os.path.join(golden, "human_cases_by_zone.png")) > 20000, "the chart is drawn")
    con.close()
    print("verify: %d checks passed from the pack alone (call %s %.3f, runner-up %s %.3f, flip %.2f per 1,000; cases "
          "%d; trap-nights run %d, lost %d; %d re-imported copies)" % (
              n[0], call, idx[call], second, idx[second], flip, sum(cases.values()), sum(run.values()),
              sum(lost.values()), len(copies)))


if __name__ == "__main__":
    main(*(sys.argv[1:3]))
