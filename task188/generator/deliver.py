"""The goldens, every figure computed from the shipped pack: the aerial spray recommendation (PDF), the zone season
workbook (XLSX) and the chart of locally acquired human cases by week of onset (PNG)."""
import datetime as dt
import os
import re
import sqlite3

import analysis as A
import world as W

MEMO = "aerial_spray_recommendation.pdf"
WORKBOOK = "zone_season_2026.xlsx"
CHART = "human_cases_by_zone.png"

WEEK32 = ("2026-08-09", "2026-08-15")
SEASON = ("2026-06-01", "2026-08-15")
FAILURES = ("FAN", "CO2", "BAT", "STOLEN")
ROW1 = re.compile(r"^(RCPHL-\d{2}-\d{4}) (RC-\d{3}) (\d{2}/\d{2}/\d{4}) (Culex \w+) (\d+) (Undet\.|\d+\.\d) "
                  r"(Positive|Negative|Inconclusive)$")
ROW2 = re.compile(r"^(RCPHL-\d{2}-\d{4}) (\d{2}/\d{2}/\d{4}) (Undet\.|\d+\.\d) (Positive|Negative)$")
CODE = {"Culex tarsalis": "CXTA", "Culex pipiens": "CXPI"}


class Pack:
    """The shipped files, read the way an analyst reads them."""

    def __init__(self, target):
        import build as B
        f = lambda k: os.path.join(target, B.FILES[k])
        con = sqlite3.connect(f("db"))
        self.sites = {r[0]: dict(trap_type=r[1], zone=r[2], section=r[3]) for r in
                      con.execute("SELECT site_id, trap_type, zone, section FROM sites")}
        self.tn = {r[0]: dict(site=r[1], set=r[2], pickup=r[3], status=r[4], entered=r[5]) for r in
                   con.execute("SELECT tn_id, site_id, set_date, pickup_date, status, entered FROM trap_nights")}
        self.catches = {}
        for t, sp, sex, n in con.execute("SELECT tn_id, species, sex, count FROM catches"):
            self.catches.setdefault(t, []).append((sp, sex, n))
        self.pools = {r[0]: dict(tn=r[1], species=r[2], females=r[3], lab=r[4]) for r in
                      con.execute("SELECT pool_id, tn_id, species, females, lab FROM pools")}
        self.first, self.final = {}, {}
        for pid, no, res in con.execute("SELECT pool_id, test_no, result FROM pool_tests ORDER BY pool_id, test_no"):
            if no == 1:
                self.first[pid] = res
            self.final[pid] = res
        con.close()
        self.county = []
        retest = {}
        import pdfplumber
        with pdfplumber.open(f("county")) as pdf:
            for page in pdf.pages:
                for line in (page.extract_text() or "").splitlines():
                    m = ROW1.match(line.strip())
                    if m:
                        self.county.append(dict(pool=m.group(1), site=m.group(2), collected=m.group(3),
                                                species=CODE[m.group(4)], females=int(m.group(5)), result=m.group(7)))
                        continue
                    m = ROW2.match(line.strip())
                    if m:
                        retest[m.group(1)] = m.group(4)
        word = {"Positive": "POS", "Negative": "NEG", "Inconclusive": "INC"}
        for r in self.county:
            self.first[r["pool"]] = word[r["result"]]
            self.final[r["pool"]] = word[retest[r["pool"]]] if r["result"] == "Inconclusive" else word[r["result"]]
        from openpyxl import load_workbook
        wb = load_workbook(f("register"), read_only=True)
        self.zone_of_section = {r[0]: r[4] for r in wb["Zone sections"].iter_rows(min_row=2, values_only=True)}
        wb = load_workbook(f("cases"), read_only=True)
        rows = list(wb["Cases"].iter_rows(values_only=True))
        head = rows[0]
        self.cases = [dict(zip(head, r)) for r in rows[1:]]

    # the week's CO2-trap females and the pools tested
    def week32_tn(self, zone, trap_type="CO2", ok=True):
        return [t for t, r in self.tn.items() if WEEK32[0] <= r["set"] <= WEEK32[1] and
                self.sites[r["site"]]["zone"] == zone and self.sites[r["site"]]["trap_type"] == trap_type and
                (r["status"] == "OK" or not ok)]

    def abundance(self, zone):
        tns = self.week32_tn(zone)
        ab = {sp: sum(n for t in tns for s, sex, n in self.catches.get(t, []) if s == sp and sex == "F") / len(tns)
              for sp in W.VECTORS}
        return ab, len(tns)

    def week32_pools(self, zone, which="CO2", result="final"):
        res = self.final if result == "final" else {k: ("POS" if v in ("POS", "INC") else "NEG")
                                                    for k, v in self.first.items()}
        out = {sp: [] for sp in W.VECTORS}
        for pid, p in self.pools.items():
            r = self.tn[p["tn"]]
            s = self.sites[r["site"]]
            if not (WEEK32[0] <= r["set"] <= WEEK32[1]) or s["zone"] != zone:
                continue
            if which == "CO2" and s["trap_type"] != "CO2":
                continue
            out[p["species"]].append((p["females"], res[pid] == "POS"))
        return out

    def index(self, zone, est="MLE", which="CO2", genus=False, result="final"):
        f = {"MLE": A.mle, "BC": A.bcmle, "MIR": A.mir}[est]
        ab, _ = self.abundance(zone)
        pl = self.week32_pools(zone, which, result)
        if genus:
            allp = pl["CXTA"] + pl["CXPI"]
            return sum(ab.values()) * f(allp)
        return sum(ab[sp] * f(pl[sp]) for sp in W.VECTORS)

    # ask B: the season's CO2 trap-nights
    def copies(self):
        """Trap-nights brought in a second time by the 22 June re-import: entered that day, dated a day after a
        night of the same site with the same status and the same catch."""
        by = {(r["site"], r["set"]): t for t, r in self.tn.items() if not r["entered"].startswith("2026-06-22")}
        out = set()
        for t, r in self.tn.items():
            if not r["entered"].startswith("2026-06-22"):
                continue
            prev = (dt.date.fromisoformat(r["set"]) - dt.timedelta(days=1)).isoformat()
            o = by.get((r["site"], prev))
            if o is not None and self.tn[o]["status"] == r["status"] and \
                    sorted(self.catches.get(o, [])) == sorted(self.catches.get(t, [])):
                out.add(t)
        return out

    def season_trap_nights(self, drop=None):
        drop = self.copies() if drop is None else drop
        out = {z: [0, 0] for z in W.ZONES}
        for t, r in self.tn.items():
            if t in drop or not (SEASON[0] <= r["set"] <= SEASON[1]):
                continue
            s = self.sites[r["site"]]
            if s["trap_type"] != "CO2":
                continue
            if r["status"] == "OK":
                out[s["zone"]][0] += 1
            elif r["status"] in FAILURES:
                out[s["zone"]][1] += 1
        return out

    # ask A: the season's locally acquired cases
    def season_cases(self):
        seen = set()
        out = {z: [] for z in W.ZONES}
        lo, hi = dt.date(2026, 6, 1), dt.date(2026, 8, 15)
        for r in self.cases:
            if r["Clinical presentation"] == "Asymptomatic blood donor" or r["Exposure"] != "Rindal County" or \
                    r["Classification"] == "Not a case":
                continue
            on = r["Onset date"]
            if on is None:
                continue
            on = on.date() if hasattr(on, "date") else on
            if not (lo <= on <= hi) or r["State case ID"] in seen:
                continue
            seen.add(r["State case ID"])
            out[self.zone_of_section[r["Residence section"]]].append(on)
        return out


def mmwr(d):
    return W.mmwr_week(d)


def figures(target):
    pk = Pack(target)
    idx = {z: pk.index(z) for z in W.ZONES}
    order = sorted(W.ZONES, key=lambda z: -idx[z])
    call, second = order[0], order[1]
    ab_c, n_c = pk.abundance(call)
    ab_s, n_s = pk.abundance(second)
    pl_c = pk.week32_pools(call)
    pl_s = pk.week32_pools(second)
    culex_c = sum(ab_c.values())
    flip = idx[second] / culex_c * 1000.0
    cases = pk.season_cases()
    tnr = pk.season_trap_nights()
    weeks = list(range(22, 33))
    by_week = {z: [sum(1 for d in cases[z] if mmwr(d) == w) for w in weeks] for z in W.ZONES}

    def side(z, ab, n, pl):
        return {"zone": z, "trap_nights": n, "tarsalis_per_tn": round(ab["CXTA"], 1), "pipiens_per_tn": round(ab["CXPI"], 1),
                "culex_per_tn": round(sum(ab.values()), 1),
                "tarsalis_pools": len(pl["CXTA"]), "tarsalis_positive": sum(x for _, x in pl["CXTA"]),
                "pipiens_pools": len(pl["CXPI"]), "pipiens_positive": sum(x for _, x in pl["CXPI"]),
                "tarsalis_mle_per_1000": round(A.mle(pl["CXTA"]) * 1000, 1),
                "pipiens_mle_per_1000": round(A.mle(pl["CXPI"]) * 1000, 1),
                "tarsalis_mir_per_1000": round(A.mir(pl["CXTA"]) * 1000, 1),
                "index": round(idx[z], 1), "index_exact": round(idx[z], 4)}
    ladder = {}
    for key, est, which in (("default", "MIR", "ALL"), ("co2_mir", "MIR", "CO2"), ("all_mle", "MLE", "ALL"),
                            ("co2_mle", "MLE", "CO2")):
        v = {z: pk.index(z, est, which) for z in W.ZONES}
        o = sorted(W.ZONES, key=lambda z: -v[z])
        ladder[key] = {"pick": o[0], "pick_figure": round(v[o[0]], 2), "second": o[1],
                       "second_figure": round(v[o[1]], 2), "margin": round(v[o[0]] / v[o[1]], 3),
                       "westbrook_place": o.index("Westbrook") + 1, "westbrook_figure": round(v["Westbrook"], 2)}
    fig = {
        "call": call, "call_figure": round(idx[call], 1), "call_figure_exact": round(idx[call], 4),
        "runner_up": second, "runner_up_figure": round(idx[second], 1), "runner_up_figure_exact": round(idx[second], 4),
        "flip_rate_per_1000": int(round(flip)), "flip_rate_exact": round(flip, 3),
        "call_rate_per_1000": round(idx[call] / culex_c * 1000.0, 1),
        "call_build": side(call, ab_c, n_c, pl_c), "runner_up_build": side(second, ab_s, n_s, pl_s),
        "bias_corrected": {"call": round(pk.index(call, "BC"), 1), "runner_up": round(pk.index(second, "BC"), 1),
                           "flip": int(round(pk.index(second, "BC") / culex_c * 1000.0))},
        "ladder": ladder,
        "cases_by_zone": {z: len(cases[z]) for z in W.ZONES}, "cases_total": sum(len(v) for v in cases.values()),
        "weeks": weeks, "cases_by_week": by_week,
        "co2_trap_nights_run": {z: tnr[z][0] for z in W.ZONES}, "co2_trap_nights_lost": {z: tnr[z][1] for z in W.ZONES},
        "reimport_copies": len(pk.copies()),
        "week": {"mmwr": 32, "from": WEEK32[0], "to": WEEK32[1]}, "flight": W.FLIGHT.isoformat(),
        "rate_basis": 1000, "pool_max": dict(W.POOL_MAX), "flight_text": "Thursday 20 August 2026",
        "protocol_clauses": {"trap_night": "3.4", "pooling": "4.1", "retest": "4.3"},
    }
    return fig


def write_chart(path, fig):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 8.5, "svg.hashsalt": "188"})
    weeks = fig["weeks"]
    zones = sorted(W.ZONES, key=lambda z: (-fig["cases_by_zone"][z], W.ZONES.index(z)))
    fig_, axes = plt.subplots(2, 4, figsize=(11.0, 5.2), sharex=True, sharey=True, dpi=110)
    top = max(max(v) for v in fig["cases_by_week"].values())
    for ax, z in zip(axes.flat, zones):
        vals = fig["cases_by_week"][z]
        ax.bar(weeks, vals, width=0.72, color="#b5532b" if z == fig["call"] else "#3b6e8f")
        for w, v in zip(weeks, vals):
            if v:
                ax.text(w, v + 0.06, str(v), ha="center", va="bottom", fontsize=7.5, color="#1c2430")
        ax.set_title("%s (%d)%s" % (z, fig["cases_by_zone"][z], ", block for 20 August" if z == fig["call"] else ""),
                     fontsize=9.2, loc="left", color="#b5532b" if z == fig["call"] else "#1c2430")
        ax.set_ylim(0, top + 1)
        ax.set_xticks(weeks)
        ax.set_xticklabels([str(w) for w in weeks], fontsize=7.2)
        ax.yaxis.get_major_locator().set_params(integer=True)
        ax.spines[["top", "right"]].set_visible(False)
        ax.grid(axis="y", color="#dfe3e8", linewidth=0.6)
        ax.set_axisbelow(True)
    for ax in axes[1]:
        ax.set_xlabel("MMWR week of onset (2026)")
    for ax in axes[:, 0]:
        ax.set_ylabel("Cases")
    first = {z: min(w for w, v in zip(weeks, fig["cases_by_week"][z]) if v) for z in W.ZONES if fig["cases_by_zone"][z]}
    last = {z: max(w for w, v in zip(weeks, fig["cases_by_week"][z]) if v) for z in W.ZONES if fig["cases_by_zone"][z]}
    top_z = zones[0]
    fig_.suptitle("%s had the season's most locally acquired cases, all with onset in weeks %d to %d; %s's %d fell ill "
                  "in week %d" % (top_z, first[top_z], last[top_z], fig["call"], fig["cases_by_zone"][fig["call"]],
                                  last[fig["call"]]) if first[fig["call"]] == last[fig["call"]] else
                  "%s had the season's most locally acquired cases" % top_z, fontsize=10.4, x=0.01, ha="left",
                  fontweight="bold")
    fig_.text(0.01, 0.925, "West Nile virus cases by MMWR week of onset, 1 June to 15 August 2026, by treatment zone as "
                           "redrawn on 1 May 2026; zone total in brackets", fontsize=8.6, color="#5b6470")
    fig_.text(0.01, 0.005, "Source: Rindal County Public Health case log of 14 August 2026, residence section on the "
                           "2026 zone map; blood donors, travel-acquired, ruled-out and repeat notifications excluded.",
              fontsize=7.2, color="#5b6470")
    fig_.tight_layout(rect=(0, 0.03, 1, 0.91))
    fig_.savefig(path, dpi=110, metadata={"Software": None})
    plt.close(fig_)


def write_workbook(path, fig):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Zones"
    ws.append(["Zone", "Locally acquired human cases, onset 1 Jun to 15 Aug 2026",
               "CO2 trap-nights run, 1 Jun to 15 Aug 2026", "CO2 trap-nights lost to trap failures, 1 Jun to 15 Aug 2026"])
    for z in W.ZONES:
        ws.append([z, fig["cases_by_zone"][z], fig["co2_trap_nights_run"][z], fig["co2_trap_nights_lost"][z]])
    ws.append(["All zones", fig["cases_total"], sum(fig["co2_trap_nights_run"].values()),
               sum(fig["co2_trap_nights_lost"].values())])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.row_dimensions[1].height = 44
    for col, w in zip("ABCD", (18, 22, 18, 22)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B2"
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    wk = wb.create_sheet("Cases by week")
    wk.append(["Zone"] + ["Week %d" % w for w in fig["weeks"]] + ["Total"])
    for z in W.ZONES:
        wk.append([z] + fig["cases_by_week"][z] + [fig["cases_by_zone"][z]])
    for c in wk[1]:
        c.font = Font(bold=True)
    wk.column_dimensions["A"].width = 18
    nt = wb.create_sheet("Notes")
    notes = [
        "Cases: Rindal County Public Health case log of 14 August 2026. Symptomatic cases exposed in Rindal County with "
        "onset from 1 June to 15 August 2026, one per state case ID. Each case is placed in the zone of its residence "
        "section on the zone map in force from 1 May 2026 (the log's zone column follows the March 2025 zone list). "
        "Presumptive viremic blood donors and travel-acquired cases are not counted.",
        "Week of onset: MMWR week, Sunday to Saturday.",
        "Trap-nights: CO2 light traps set from 1 June to 15 August 2026 (Monday and Tuesday nights, and Maple Ridge's "
        "enhanced Wednesday and Thursday nights of 17 to 25 June). Run: status OK. Lost to trap failures: fan, CO2, "
        "battery and stolen. Traps not set (NS) are neither. The %d records the second load of 22 June entered "
        "again (dated the pickup morning, with the same site, status and catch as the night already in the "
        "database) are counted once." % fig["reimport_copies"],
    ]
    notes.insert(0, "As of the field database extract of 17 August 2026, 08:20, and the county case log of 14 August "
                    "2026. Zones as redrawn on 1 May 2026.")
    for n in notes:
        nt.append([n])
        nt.cell(row=nt.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    nt.column_dimensions["A"].width = 120
    wb.save(path)


def write_memo(path, fig):
    from reportlab import rl_config
    rl_config.invariant = 1
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    import docs as D   # registers the fonts
    ink, grey = colors.HexColor("#1c2430"), colors.HexColor("#5b6470")
    body = ParagraphStyle("b", fontName="Sans", fontSize=9.8, leading=13.6, textColor=ink, spaceAfter=6)
    head = ParagraphStyle("h", fontName="Sans-B", fontSize=13.5, leading=17, textColor=ink, spaceAfter=4)
    h2 = ParagraphStyle("h2", fontName="Sans-B", fontSize=10.2, leading=13.5, textColor=ink, spaceBefore=6, spaceAfter=3)
    meta = ParagraphStyle("m", fontName="Sans", fontSize=8.4, leading=11, textColor=grey, spaceAfter=6)
    cell = ParagraphStyle("c", fontName="Sans", fontSize=8.6, leading=10.6, textColor=ink)
    cellb = ParagraphStyle("cb", parent=cell, fontName="Sans-B")
    c, s = fig["call_build"], fig["runner_up_build"]
    who, role = W.PEOPLE["surv"]
    boss, brole = W.PEOPLE["ops"]
    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.9 * inch, rightMargin=0.9 * inch, topMargin=0.8 * inch,
                            bottomMargin=0.8 * inch, title="Aerial block for Thursday 20 August 2026", author=who,
                            subject="", creator="", producer="")
    x = [Paragraph(W.DISTRICT.upper() + " · OPERATIONS", meta),
         Paragraph("Aerial block for Thursday 20 August 2026: %s" % fig["call"], head),
         Paragraph("To %s, %s. From %s, %s. Monday 17 August 2026, on the surveillance of MMWR week 32 (trap-nights "
                   "of 10 and 11 August)." % (boss, brole, who, role), meta),
         Paragraph("<b>Fly %s.</b> It had the most West Nile virus-infected Culex females per CO2 trap-night last week: "
                   "an estimated <b>%.1f</b>. The runner-up is <b>%s</b> at <b>%.1f</b>. %s keeps the flight unless "
                   "the infection of its CO2-caught Culex females is below <b>%d per 1,000</b>; last week it was %.1f "
                   "per 1,000." % (fig["call"], fig["call_figure"], fig["runner_up"], fig["runner_up_figure"],
                                    fig["call"], fig["flip_rate_per_1000"], fig["call_rate_per_1000"]), body),
         Paragraph("How the figure is built", h2),
         Paragraph("For each zone, the Culex tarsalis and Culex pipiens females caught per CO2 trap-night that ran "
                   "(failed traps left out), times the share of those same females infected, summed over the two "
                   "species. The infected share comes from the pools made from the CO2 traps' own catch, estimated "
                   "by maximum likelihood for pools of varying size. The bias-corrected estimate gives the same "
                   "figures to one decimal.", body),
         Paragraph("Two shortcuts give the wrong zone. Counting each positive pool as one infected mosquito (the "
                   "minimum infection rate) understates our laboratory's pools of 50: %d of %s's %d Culex tarsalis "
                   "pools were positive, and a pool that full of positives holds more than one infected female. And "
                   "the gravid traps in Old Town and Riverside South catch females that have already fed; their "
                   "pools run several times hotter than the host-seeking females the CO2 traps catch, so they cannot "
                   "stand in for them. Coulee Hills has the most human cases this season, but they fell ill between late "
                   "June and late July, in the coulee focus that was flown on 16 and 23 July." % (
                       c["tarsalis_positive"], fig["call"], c["tarsalis_pools"]), body),
         Paragraph("The two leading zones, week 32", h2)]
    rows = [["", fig["call"], fig["runner_up"]],
            ["CO2 trap-nights run", str(c["trap_nights"]), str(s["trap_nights"])],
            ["Culex tarsalis females per trap-night", "%.1f" % c["tarsalis_per_tn"], "%.1f" % s["tarsalis_per_tn"]],
            ["Culex pipiens females per trap-night", "%.1f" % c["pipiens_per_tn"], "%.1f" % s["pipiens_per_tn"]],
            ["Culex tarsalis pools (positive)", "%d (%d)" % (c["tarsalis_pools"], c["tarsalis_positive"]),
             "%d (%d)" % (s["tarsalis_pools"], s["tarsalis_positive"])],
            ["Culex pipiens pools (positive)", "%d (%d)" % (c["pipiens_pools"], c["pipiens_positive"]),
             "%d (%d)" % (s["pipiens_pools"], s["pipiens_positive"])],
            ["Culex tarsalis infected per 1,000", "%.1f" % c["tarsalis_mle_per_1000"], "%.1f" % s["tarsalis_mle_per_1000"]],
            ["Culex pipiens infected per 1,000", "%.1f" % c["pipiens_mle_per_1000"], "%.1f" % s["pipiens_mle_per_1000"]],
            ["Infected Culex females per CO2 trap-night", "%.1f" % fig["call_figure"], "%.1f" % fig["runner_up_figure"]]]
    t = Table([[Paragraph(v, cellb if (i == 0 or i == len(rows) - 1) else cell) for v in r] for i, r in enumerate(rows)],
              colWidths=[3.1 * inch, 1.5 * inch, 1.5 * inch], hAlign="LEFT")
    t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, 0), 0.6, ink), ("LINEABOVE", (0, -1), (-1, -1), 0.6, ink),
                           ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
    x += [t, Paragraph("Source: field database extract of 17 August 2026, 08:20 (CO2 trap-nights, catches and the "
                       "district laboratory's results) and the county laboratory's report of 14 August 2026, whose week-32 "
                       "results are not yet keyed into the database. Infection estimated by maximum likelihood from each "
                       "zone's CO2-trap pools, by species.", meta), Spacer(1, 6),
          Paragraph("The flip point is %s's figure over %s's %.1f Culex females per CO2 trap-night. The season's human "
                    "cases and trap-nights by zone are in zone_season_2026.xlsx, and the cases by week of onset in "
                    "human_cases_by_zone.png." % (fig["runner_up"], fig["call"], c["culex_per_tn"]), body),
          Paragraph("Neither zone lost a trap-night to a failure in week 32 (protocol 3.4), and neither had an "
                    "inconclusive pool. Designation to %s by 17:00 today (agreement clause 2.2)." % W.APPLICATOR, meta)]

    def foot(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Sans", 7.2)
        canvas.setFillColor(grey)
        canvas.drawString(0.9 * inch, 0.5 * inch, "%s | Operations | Aerial block, MMWR week 33" % W.DISTRICT_SHORT)
        canvas.drawRightString(letter[0] - 0.9 * inch, 0.5 * inch, "Page %d" % doc_.page)
        canvas.restoreState()
    doc.build(x, onFirstPage=foot, onLaterPages=foot)
