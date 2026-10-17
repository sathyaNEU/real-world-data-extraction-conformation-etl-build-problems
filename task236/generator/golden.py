"""The golden deliverables: influenza_testing_break_flags_week36_2027.xlsx and testing_change_breaks_week36_2027.png."""
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import params as P
import postprocess

XLSX = "influenza_testing_break_flags_week36_2027.xlsx"
PNG = "testing_change_breaks_week36_2027.png"
RED = "#B03A2E"
GREY = "#8EA3B9"
INK = "#2F3B45"
HEAD = ["Region", "Region code", "Register entry", "Testing change", "In force from", "Report week of change",
        "Testing-change break flag", "Notified cases, season to date (weeks 14 to 36)",
        "Notified rate per 100,000 residents, week 36"]


def workbook(path, A):
    wb = Workbook(); ws = wb.active; ws.title = "Flags week 36"
    ws.append(["Week 36 of 2027 (6 to 12 September): testing-change break flags on four notified series, %s" %
               ", ".join(r["name"] for r in A["regions"] if r["flag"])])
    ws.append(["Signed off Tuesday 14 September 2027 under CDB-FLU-03 s.4. Counts as the week 36 report counts them: "
               "cases notified by 09:00 on Monday 13 September 2027."])
    ws.append([])
    ws.append(HEAD)
    hr = ws.max_row
    for r in A["regions"]:
        ws.append([r["name"], r["code"], r["entry"], r["change"], r["date"], r["week"], "Yes" if r["flag"] else "No",
                   r["season"], r["rate36"]])
    last = ws.max_row
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"].font = Font(italic=True, size=9, color="555555")
    for c in ws[hr]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="E8EEF0")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=hr + 1, max_row=last):
        row[3].alignment = Alignment(wrap_text=True, vertical="top")
        row[4].number_format = "d mmm yyyy"
        row[7].number_format = "#,##0"
        row[8].number_format = "0.0"
        if row[6].value == "Yes":
            row[6].font = Font(bold=True, color=RED.strip("#"))
    for i, w in enumerate([15, 9, 12, 58, 13, 11, 12, 16, 15], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[hr].height = 44
    ws.freeze_panes = ws.cell(row=hr + 1, column=3)
    ws.print_title_rows = "%d:%d" % (hr, hr)
    ws.page_setup.orientation = "landscape"
    nt = wb.create_sheet("Notes")
    for line in A["notes"]:
        nt.append([line])
    nt.column_dimensions["A"].width = 125
    for row in nt.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    postprocess.scrub(path, P.PEOPLE["epi"], "2027-09-14")
    postprocess.finish_office(path, "2027-09-13T23:10:00Z", "2027-09-14T01:25:00Z", P.DEPT)


def chart(path, A):
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 8.5})
    fig, axes = plt.subplots(2, 4, figsize=(13.0, 6.6), dpi=150, sharex=True)
    for ax, r in zip(axes.flat, A["regions"]):
        pnl = r["panel"]
        wk, n, x = pnl["weeks"], pnl["notified"], pnl["index"]
        col = [RED if (r["flag"] and w >= r["week"]) else GREY for w in wk]
        ax.bar(wk, n, color=col, width=0.72, zorder=3)
        ax.bar([wk[-1]], [n[-1]], color="none", edgecolor=INK, hatch="///", width=0.72, linewidth=0.0, zorder=4)
        ax.axvline(r["week"] - 0.5, color=INK, linestyle="--", linewidth=1.0, zorder=5)
        top = max(n) * 1.18
        ax.set_ylim(0, top)
        ax.text(r["week"] - 0.35, top * 0.97, "change wk %d" % r["week"], fontsize=7.2, color=INK, va="top")
        ax2 = ax.twinx()
        ax2.plot(wk, x, color="#1F6E8C", linewidth=1.4, marker="o", markersize=2.5, zorder=6)
        ax2.set_ylim(0, max(v for v in x) * 1.25)
        ax2.tick_params(axis="y", labelsize=7, colors="#1F6E8C")
        if ax in axes[:, -1]:
            ax2.set_ylabel("influenza activity", fontsize=8, color="#1F6E8C")
        for s in ("top",):
            ax.spines[s].set_visible(False); ax2.spines[s].set_visible(False)
        title = "%s: %s" % (r["name"], "FLAG from week %d" % r["week"] if r["flag"] else "no flag")
        ax.set_title(title, fontsize=9.5, fontweight="bold", color=RED if r["flag"] else INK, loc="left", pad=15)
        ax.text(0.0, 1.02, r["kind_short"], transform=ax.transAxes, fontsize=7.4, color="#444444", va="bottom")
        ax.grid(axis="y", color="#E8E8E8", zorder=0)
        ax.tick_params(axis="y", labelsize=7)
        ax.set_xticks(wk[::2])
    for ax in axes[:, 0]:
        ax.set_ylabel("notified cases", fontsize=8)
    for ax in axes[1, :]:
        ax.set_xlabel("report week, 2027", fontsize=8)
    fig.tight_layout(rect=(0, 0, 1, 0.87), h_pad=2.2)
    fl = [r["name"] for r in A["regions"] if r["flag"]]
    fig.text(0.012, 0.985, "Flag %s and %s from their change weeks; the other four series stay comparable"
             % (", ".join(fl[:-1]), fl[-1]), fontsize=12.5, fontweight="bold", va="top", color=INK)
    fig.text(0.012, 0.945, "Bars: notified influenza cases by report week (week 36 hatched: as counted at 09:00 Monday 13 "
                           "September). Line (right axis): sentinel general practice influenza-like-illness consultations "
                           "per 100,000 patients\ntimes the share of protocol swabs positive for influenza. Red bars: weeks "
                           "a flag covers. Dashed line: the testing change.", fontsize=8.2, va="top", color="#333333",
             linespacing=1.4)
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)


def write(golden_dir, A):
    os.makedirs(golden_dir, exist_ok=True)
    workbook(os.path.join(golden_dir, XLSX), A)
    chart(os.path.join(golden_dir, PNG), A)
