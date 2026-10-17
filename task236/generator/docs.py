"""task236 documents: the weekly influenza report's data quality protocol and the sentinel general practice network's
protocol (pdf), the testing change register (xlsx), the reporting services list (csv), the data dictionary and the
readme (md)."""
import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess

NAMES = {r[0]: r[1] for r in P.REGIONS}


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=13, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=14, leading=18, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=14, spaceBefore=8,
                        spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10, textColor=colors.HexColor("#444444"))
    return dict(body=body, h1=h1, h2=h2, small=small)


def _footer(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(18 * mm, 10 * mm, text)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _table(rows, widths, size=8.0):
    t = Table(rows, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", size), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", size),
                           ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF0")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


# ------------------------------------------------------------------ the report's data quality protocol
def protocol(path):
    S = _styles(); b, h1, h2, sm = S["body"], S["h1"], S["h2"], S["small"]
    st = [Paragraph("Weekly influenza report: data quality protocol", h1),
          Paragraph("%s, %s. Protocol CDB-FLU-03, version 4, in force from 1 March 2027 (replaces version 3 of April "
                    "2026). Owner: %s, surveillance epidemiologist." % (P.BRANCH, P.DEPT, P.PEOPLE["epi"]), sm),
          Spacer(1, 4),
          Paragraph("1. Purpose", h2),
          Paragraph("From April to October the branch publishes a weekly influenza report: notified influenza by health "
                    "region and report week, rates per 100,000 residents, the season to date, and a sentinel section "
                    "from the sentinel general practice network (protocol SGP-2027). This protocol sets how the notified "
                    "series are counted, the cut-off each report is counted at, and the flags the report carries where a "
                    "series is not comparable across a change.", b),
          Paragraph("2. Notified cases", h2),
          Paragraph("Influenza is notified through the notification gateway, which receives every laboratory's and every "
                    "point-of-care service's influenza results. A case is a person with influenza A or B detected by a "
                    "laboratory (a confirmed case) or by a point-of-care antigen test alone (a probable case); the "
                    "notified series count confirmed and probable cases. Positive results from specimens collected "
                    "within 14 days of the collection of a case's first positive specimen belong to that case; a "
                    "positive specimen collected 14 or more days after it starts a new case. A case is counted in the "
                    "report week of its first positive specimen's collection date. Report weeks are ISO weeks, Monday to "
                    "Sunday, local time.", b),
          Paragraph("A case is counted in the health region of its residential postcode under the postcode concordance in "
                    "force at the report. A case whose residential postcode was not recorded is counted in the region of "
                    "the requesting service's postcode. Residents of other states are not counted.", b),
          Paragraph("A case is notified when the gateway receives the first result reporting influenza for it. Each "
                    "week's report counts the cases notified by 09:00 on the Monday after its report week; a case "
                    "notified later is counted from the next report. The season to date runs from report week 14.", b),
          Paragraph("3. Rates", h2),
          Paragraph("Rates are cases per 100,000 residents of the region: the estimated resident population at 30 June "
                    "2026 of the postcodes in the region under the concordance in force at the report.", b),
          Paragraph("4. Testing-change break flags", h2),
          Paragraph("A region's notified series depends on who is tested as well as on how much influenza there is. The "
                    "laboratory liaison keeps the testing change register: every change in a region's testing that the "
                    "laboratories, the point-of-care program or the local health districts notify to the branch.", b),
          Paragraph("When the register records a change in a region's testing, the region's notified series carries a "
                    "testing-change break flag from the report week of the change if the change raised or lowered the "
                    "share of the region's influenza infections that are notified by more than a quarter of its level "
                    "before the change, comparing the four report weeks from the week of the change with the four report "
                    "weeks before it. Infections count whether or not the person sought care or was tested.", b),
          Paragraph("A flagged series is published with a break marker at the week of the change, and comparisons across "
                    "the marker are footnoted in the report. A series that is not flagged is published without a marker. "
                    "The surveillance epidemiologist signs off the flags on the Tuesday after the report week; a flag, "
                    "once published, stays on the series for the rest of the season.", b),
          Paragraph("5. Weekly checks", h2),
          _table([["Check", "Run on", "Target", "Action"],
                  ["Residential postcode recorded", "each service's results", "98 per cent or more",
                   "laboratory liaison follows up the service"],
                  ["Results received by the cut-off", "the report week's specimens", "reviewed weekly",
                   "late services named in the report's notes"],
                  ["Gateway message errors", "rejected messages", "none outstanding at the cut-off",
                   "resubmission requested"],
                  ["Testing change register", "new entries", "every entry reviewed before sign-off",
                   "flag assessed under section 4"]],
                 [44 * mm, 40 * mm, 38 * mm, 52 * mm]),
          Spacer(1, 6),
          Paragraph("6. Roles", h2),
          _table([["Role", "Holder", "In this protocol"],
                  ["Surveillance epidemiologist", P.PEOPLE["epi"], "counts the series and signs off the flags (s.4)"],
                  ["Laboratory liaison", P.PEOPLE["lab"], "keeps the testing change register (s.4); weekly checks (s.5)"],
                  ["Sentinel network coordinator", P.PEOPLE["sentinel"], "the report's sentinel section (SGP-2027)"]],
                 [52 * mm, 36 * mm, 86 * mm]),
          ]
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title="Weekly influenza report: data quality protocol",
                            author=P.PEOPLE["epi"], creator=P.PEOPLE["epi"], invariant=1)
    foot = "CDB-FLU-03 v4 | %s | %s" % (P.BRANCH, P.DEPT)
    doc.build(st, onFirstPage=_footer(foot), onLaterPages=_footer(foot))
    postprocess.scrub(path, P.PEOPLE["epi"], "2027-02-22")


# ------------------------------------------------------------------ the sentinel network's protocol
def sentinel_protocol(path):
    S = _styles(); b, h1, h2, sm = S["body"], S["h1"], S["h2"], S["small"]
    rows = [["Region", "Practices"]] + [["%s (%s)" % (NAMES[c], c), str(P.SENTINEL_PRACTICES[c])] for c in NAMES]
    rows.append(["State", str(sum(P.SENTINEL_PRACTICES.values()))])
    st = [Paragraph("Sentinel general practice influenza surveillance: protocol for 2027", h1),
          Paragraph("%s, %s. Protocol SGP-2027, for the season from 5 April 2027 (week 14). Coordinator: %s." %
                    (P.BRANCH, P.DEPT, P.PEOPLE["sentinel"]), sm),
          Spacer(1, 4),
          Paragraph("1. The network", h2),
          Paragraph("The network is a fixed panel of general practices recruited in every health region, which take part "
                    "for the whole season. Each practice's list size is the number of patients registered with it, taken "
                    "from the practice's records at the start of each quarter and reported with each weekly return.", b),
          _table(rows, [70 * mm, 30 * mm]),
          Spacer(1, 6),
          Paragraph("2. Influenza-like illness", h2),
          Paragraph("A patient with influenza-like illness (ILI) has a fever (measured or reported), a cough and fatigue, "
                    "with onset in the past seven days. The definition is clinical; the practice records a consultation as "
                    "ILI whether or not the patient is swabbed or tested elsewhere.", b),
          Paragraph("3. Weekly returns", h2),
          Paragraph("By Wednesday each practice returns, for the week Monday to Sunday just ended, its consultations (all "
                    "patients seen by its participating general practitioners) and its ILI consultations. A practice "
                    "that misses a return is shown as not received for that week; the coordinator does not fill in "
                    "missing returns. The network's ILI rate for a region and week is the ILI consultations per 100,000 "
                    "patients on the lists of the region's practices that made a return.", b),
          Paragraph("4. Protocol swabs", h2),
          Paragraph("Each practice swabs the first two patients with ILI on each day it is open, up to ten in a week, with "
                    "the patient's consent. Swabs go by courier to the Public Health Laboratory, which tests them with a "
                    "multiplex PCR for influenza A, influenza B, respiratory syncytial virus (RSV) and SARS-CoV-2. "
                    "Protocol swabs are surveillance specimens: results go to the coordinator without the patient's "
                    "identity and are not notified, and the practice is not sent a result for the patient. A swab that "
                    "cannot be tested is recorded with the laboratory's comment and no result. The network's swab "
                    "positivity for a region and week is the share of its tested protocol swabs with influenza A or B "
                    "detected.", b),
          Paragraph("5. Publication", h2),
          Paragraph("The coordinator publishes each region's ILI rate and swab positivity in the sentinel section of the "
                    "weekly influenza report. Practices are not named.", b),
          ]
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title="Sentinel general practice influenza surveillance: protocol 2027",
                            author=P.PEOPLE["sentinel"], creator=P.PEOPLE["sentinel"], invariant=1)
    foot = "SGP-2027 | %s | %s" % (P.BRANCH, P.DEPT)
    doc.build(st, onFirstPage=_footer(foot), onLaterPages=_footer(foot))
    postprocess.scrub(path, P.PEOPLE["sentinel"], "2027-03-15")


# ------------------------------------------------------------------ the testing change register
EARLIER = [
    ("TCR-27-001", "State", "State", dt.date(2027, 3, 1), "Point-of-care services onboarded to the notification gateway "
     "(rapid antigen results, positive and negative, reported from 1 March)", dt.date(2027, 2, 19)),
    ("TCR-27-002", "State", "State", dt.date(2027, 3, 29), "Public Health Laboratory: influenza PCR moves to the "
     "laboratory's new respiratory assay; subtyping of influenza A positives continues", dt.date(2027, 3, 22)),
    ("TCR-27-003", "FW", NAMES["FW"], dt.date(2027, 5, 10), "Inland Diagnostic Laboratories: twice-weekly courier from "
     "the Far West collection centres becomes daily", dt.date(2027, 5, 4)),
]


def register(path, reg):
    wb = Workbook(); ws = wb.active; ws.title = "Register"
    ws.append(["entry", "region_code", "region_name", "effective_date", "change", "recorded_by", "entered"])
    for e in EARLIER:
        ws.append([e[0], e[1], e[2], e[3], e[4], P.PEOPLE["lab"], e[5]])
    for r in reg.itertuples():
        ws.append([r.entry, r.region_code, r.region_name, dt.date.fromisoformat(r.effective_date), r.change, r.recorded_by,
                   dt.date.fromisoformat(r.entered)])
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="E8EEF0")
    for i, w in enumerate([12, 12, 15, 14, 92, 14, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        row[3].number_format = "yyyy-mm-dd"; row[6].number_format = "yyyy-mm-dd"
        row[4].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    nt = wb.create_sheet("Notes")
    for line in ("Testing change register, influenza season 2027. Kept by the laboratory liaison (%s)." % P.PEOPLE["lab"],
                 "An entry records a change in a region's testing that a laboratory, the point-of-care program or a local "
                 "health district notified to the branch, with the date it took effect.",
                 "Entries are reviewed for the report's testing-change break flags under protocol CDB-FLU-03 section 4."):
        nt.append([line])
    nt.column_dimensions["A"].width = 120
    wb.save(path)
    postprocess.scrub(path, P.PEOPLE["lab"], "2027-08-09")
    postprocess.finish_office(path, "2027-02-19T09:05:00Z", "2027-08-09T03:40:00Z", P.DEPT)


# ------------------------------------------------------------------ the reporting services
SERVICES_NOTE = {
    "CWP": "merged into Western Pathology Network on %s" % P.CHANGES["CW"][0].isoformat(),
    "IDL": "merged into Western Pathology Network on %s" % P.CHANGES["CW"][0].isoformat(),
    "WPN": "reports from %s (Coolabah West Pathology and Inland Diagnostic Laboratories)" % P.CHANGES["CW"][0].isoformat(),
    "PHL": "state reference laboratory; subtypes influenza A",
}


def services():
    import pandas as pd
    rows = []
    for code, (name, kind) in P.LABS.items():
        served = [c for c, labs in P.REGION_LABS.items() if code in labs]
        if code == "WPN":
            served = sorted({c for c, labs in P.REGION_LABS.items() if set(labs) & {"CWP", "IDL"}},
                            key=lambda c: list(NAMES).index(c))
        rows.append((code, name, kind, " ".join(served), SERVICES_NOTE.get(code, "")))
    rows.append(("POCnnnn", "Point-of-care testing services (general practices)", "point of care",
                 "where funded (see the register)", "one code per practice; reporting since 2027-03-01"))
    return pd.DataFrame(rows, columns=["service_code", "service_name", "service_type", "regions_served", "note"])


# ------------------------------------------------------------------ the dictionary and the readme
def dictionary(files):
    return """# Data dictionary

Dates are local (Australian Eastern Standard Time, UTC+10, no daylight saving from April to September) unless a field
says UTC. Report weeks are ISO weeks, Monday to Sunday.

## {res}

One row per result the notification gateway received from 5 April 2027 to the extract at 06:30 on Tuesday 14 September
2027, for specimens collected from 5 April to 12 September 2027.

| field | meaning |
|---|---|
| result_id | the gateway's id for the result message |
| specimen_id | the reporting service's accession number for the specimen |
| service_code | the reporting service (see {svc}) |
| service_type | laboratory or point of care |
| method | PCR, multiplex respiratory panel or rapid antigen |
| analyte | what the result reports on: Influenza A and B, or, for a multiplex panel, Influenza A or Influenza B on its own row |
| result | the result as the service worded it |
| result_status | P for a preliminary report, F for a final report |
| patient_key | the gateway's key for the person, the same for every service |
| sex | F, M or X |
| age_years | age at collection |
| residential_postcode | the postcode of the person's residential address as given on the request; blank, or 0000 at some services, when the request had none |
| requester_postcode | the postcode of the requesting practice, hospital or clinic |
| requester_setting | general practice, emergency department, hospital inpatient or drive-through clinic |
| collection_date | the specimen's collection date |
| received_utc | when the result reached the gateway, UTC |

## {ret}

One row per sentinel practice and report week, weeks 14 to 36 of 2027.

| field | meaning |
|---|---|
| practice_code | the network's code for the practice |
| region_code | the health region the practice is in |
| iso_week | the report week |
| week_ending | the Sunday the week ends |
| list_size | the practice's registered patients, as at the start of the quarter |
| consultations | consultations by the practice's participating general practitioners in the week; blank when the return was not received |
| ili_consultations | consultations for influenza-like illness in the week; blank when the return was not received |
| return_status | received or not received |

## {sw}

One row per protocol swab collected by a sentinel practice, weeks 14 to 36 of 2027.

| field | meaning |
|---|---|
| swab_id | the network's id for the swab |
| practice_code | the practice that took it |
| collection_date | the collection date |
| influenza_a, influenza_b, rsv, sars_cov_2 | detected or not detected by the Public Health Laboratory's multiplex PCR; blank when the swab was not tested |
| lab_comment | the laboratory's comment when the swab was not tested |

## {conc}

The postcodes of the state and their health regions, with the dates each assignment is in force (effective_to blank:
in force now).

## {pop}

Estimated resident population at 30 June 2026 by postcode (the state statistician's postcode series, released March
2027).

## {svc}

The services that report influenza results to the gateway.

## {reg}

Sheet Register: the testing change register, one row per change. Sheet Notes: about the register.

## {vac}

Influenza vaccine doses recorded in the immunisation register by region and week, general practice and pharmacy.

## {tat}

Each laboratory's weekly turnaround, specimen collection to the final report reaching the gateway, by week of collection
(collection taken as 10:00 on the collection date), weeks 14 to 35.
""".format(**files)


def readme(files):
    return """# Weekly influenza report, week 36 of 2027: extract for the testing-change break flags

Extracted at 06:30 on Tuesday 14 September 2027 for the week 36 report (the week of Monday 6 to Sunday 12 September) and
the sign-off of its testing-change break flags.

- `{prot}`: the report's data quality protocol (CDB-FLU-03 version 4).
- `{sprot}`: the sentinel general practice network's protocol for 2027.
- `{res}`: the notification gateway's influenza results.
- `{ret}`: the sentinel practices' weekly returns.
- `{sw}`: the sentinel practices' protocol swabs and their results.
- `{reg}`: the testing change register.
- `{conc}`: the postcode concordance.
- `{pop}`: estimated resident population by postcode.
- `{svc}`: the reporting services.
- `{vac}`: influenza vaccine doses by region and week, from the immunisation register.
- `{tat}`: laboratory turnaround by week, from the gateway.
- `DATA_DICTIONARY.md`: the fields of every file.
""".format(**files)
