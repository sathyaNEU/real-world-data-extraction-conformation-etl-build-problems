"""Independent verifier: reads only the shipped bundle (target/, golden/, figures.json, submission.md) and recomputes
every graded figure with its own code, then checks the goldens and the write-up carry them."""
import csv
import datetime as dt
import json
import os
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import load_workbook

RES = "gateway_influenza_results_2027-04-05_to_2027-09-12.csv"
RET = "sentinel_gp_weekly_returns_2027_wk14-36.csv"
SWB = "sentinel_gp_protocol_swabs_2027_wk14-36.csv"
CONC = "postcode_health_region_concordance.csv"
POP = "postcode_erp_30_june_2026.csv"
REG = "testing_change_register_2027.xlsx"
XLSX = "influenza_testing_break_flags_week36_2027.xlsx"
PNG = "testing_change_breaks_week36_2027.png"
CUT = "2027-09-12T23:00:00Z"            # 09:00 Monday 13 September 2027 at UTC+10
EIGHT = ["FN", "NC", "TB", "CW", "MN", "MS", "MW", "SC"]


def week(iso):
    return dt.date.fromisoformat(iso).isocalendar()[1]


def reports_influenza(text):
    t = text.lower()
    if "not detected" in t or t.startswith("negative") or t.startswith("no influenza"):
        return False
    return "detected" in t or "positive" in t


def rows(target, name):
    with open(os.path.join(target, name), newline="") as fh:
        return list(csv.DictReader(fh))


def compute(target):
    regions = {}
    for r in rows(target, CONC):
        if r["effective_to"] == "":
            regions[r["postcode"]] = r["region_code"]
    erp = defaultdict(int)
    for r in rows(target, POP):
        erp[regions[r["postcode"]]] += int(r["erp_30_june_2026"])
    # the person behind each result: a specimen is taken from one person, so gateway keys that share one are one
    # person; a key whose ages jump by more than a year carries two people (a year more is a birthday)
    res = rows(target, RES)
    parent = {}

    def find(k):
        while parent.get(k, k) != k:
            k = parent[k]
        return k
    keys_of = defaultdict(set)
    for r in res:
        keys_of[r["specimen_id"]].add(r["patient_key"])
    for ks in keys_of.values():
        ks = sorted(ks)
        for k in ks[1:]:
            a, b = find(ks[0]), find(k)
            if a != b:
                parent[max(a, b)] = min(a, b)
    ages = defaultdict(set)
    for r in res:
        ages[find(r["patient_key"])].add(int(r["age_years"]))
    part = {}
    for k, a in ages.items():
        a = sorted(a)
        lab, n = {a[0]: 0}, 0
        for x, y in zip(a, a[1:]):
            n += y - x > 1
            lab[y] = n
        part[k] = lab
    # cases: each person's positive results in collection order, a new case 14 days or more after the case's first
    by_person = defaultdict(list)
    for r in res:
        if reports_influenza(r["result"]):
            k = find(r["patient_key"])
            by_person[(k, part[k][int(r["age_years"])])].append(
                (r["collection_date"], r["received_utc"], r["residential_postcode"], r["requester_postcode"]))
    cases = []
    for key, xs in by_person.items():
        xs.sort()
        cur = None
        for d, rec, pc, req in xs:
            day = dt.date.fromisoformat(d)
            if cur is None or (day - cur["first"]).days >= 14:
                cur = {"first": day, "rec": rec, "pc": "", "req": req}
                cases.append(cur)
            elif rec < cur["rec"]:
                cur["rec"] = rec
            if not cur["pc"] and pc not in ("", "0000"):
                cur["pc"] = pc
    count = defaultdict(int)
    for c in cases:
        region = regions.get(c["pc"]) if c["pc"] else regions.get(c["req"])
        if region is None or c["rec"] > CUT:
            continue
        count[(region, c["first"].isocalendar()[1])] += 1
    # the sentinel network: ILI per 100,000 patients of the practices that returned, times the tested swabs' share
    ili, lst, home = defaultdict(float), defaultdict(float), {}
    for r in rows(target, RET):
        home[r["practice_code"]] = r["region_code"]
        if r["return_status"] == "received":
            k = (r["region_code"], int(r["iso_week"]))
            ili[k] += float(r["ili_consultations"]); lst[k] += float(r["list_size"])
    tested, flu = defaultdict(int), defaultdict(int)
    for r in rows(target, SWB):
        if r["influenza_a"] == "":
            continue
        k = (home[r["practice_code"]], week(r["collection_date"]))
        tested[k] += 1
        flu[k] += r["influenza_a"] == "detected" or r["influenza_b"] == "detected"
    act = {k: ili[k] / lst[k] * 1e5 * flu[k] / tested[k] for k in ili if tested.get(k)}
    wb = load_workbook(os.path.join(target, REG), read_only=True)
    change = {}
    for r in list(wb["Register"].iter_rows(values_only=True))[1:]:
        if r[1] in EIGHT:
            d = r[3].date() if hasattr(r[3], "date") else r[3]
            change[r[1]] = d.isocalendar()[1]
    ratio = {}
    for code, cw in change.items():
        b, a = range(cw - 4, cw), range(cw, cw + 4)
        n = lambda ws: sum(count.get((code, w), 0) for w in ws)
        x = lambda ws: sum(act[(code, w)] for w in ws)
        ratio[code] = (n(a) / x(a)) / (n(b) / x(b))
    flags = sorted(c for c, v in ratio.items() if v > 1.25 or v < 0.75)
    season = {c: sum(v for (r, w), v in count.items() if r == c and 14 <= w <= 36) for c in EIGHT}
    rate = {c: float(Decimal(repr(count.get((c, 36), 0) / erp[c] * 1e5)).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP))
            for c in EIGHT}
    before_after = {c: (sum(count.get((c, w), 0) for w in range(change[c] - 4, change[c])),
                        sum(count.get((c, w), 0) for w in range(change[c], change[c] + 4))) for c in EIGHT}
    return dict(flags=flags, ratio=ratio, season=season, rate=rate, ba=before_after, change=change)


def main(target, golden, figs_path, sub_path):
    V = compute(target)
    assert V["flags"] == ["FN", "MS", "MW", "SC"], V["flags"]
    wb = load_workbook(os.path.join(golden, XLSX))
    ws = wb["Flags week 36"]
    body = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[1]]
    assert [r[1] for r in body] == EIGHT
    for r in body:
        code = r[1]
        assert (r[6] == "Yes") == (code in V["flags"]), r
        assert r[5] == V["change"][code], r
        assert r[7] == V["season"][code], (r, V["season"][code])
        assert abs(r[8] - V["rate"][code]) < 1e-9, (r, V["rate"][code])
    png = os.path.join(golden, PNG)
    assert open(png, "rb").read(8) == b"\x89PNG\r\n\x1a\n" and os.path.getsize(png) > 40_000
    assert sorted(os.listdir(golden)) == sorted([XLSX, PNG])
    figs = json.load(open(figs_path))
    names = {r["code"]: r["region"] for r in figs["deliverables"][XLSX]}
    assert sorted(figs["call"]) == sorted(names[c] for c in V["flags"])
    for r in figs["deliverables"][XLSX]:
        c = r["code"]
        assert r["season_to_date"] == V["season"][c] and r["rate_week36"] == "%.1f" % V["rate"][c]
        assert r["flag"] == ("yes" if c in V["flags"] else "no")
        assert figs["components"]["notified_before_after"][r["region"]] == list(V["ba"][c])
        assert figs["components"]["share_notified_after_over_before"][r["region"]] == "%.2f" % round(V["ratio"][c], 2)
    sub = open(sub_path).read()
    assert "\u2014" not in sub and "\u2013" not in sub
    block4 = sub.split("## 4. Deliverable Answers")[1]
    for r in figs["deliverables"][XLSX]:
        c = r["code"]
        line = "   - %s (%s), change in report week %d: %s, %s notified cases season to date, %s per 100,000 in week 36" % (
            r["region"], c, V["change"][c], "flagged" if c in V["flags"] else "not flagged", format(V["season"][c], ","),
            "%.1f" % V["rate"][c])
        assert line in block4, line
    block2 = sub.split("## 2. Critical Components")[1].split("## 3.")[0]
    for c in EIGHT:
        b, a = V["ba"][c]
        assert "%s **%s** and **%s**" % (names[c], format(b, ","), format(a, ",")) in block2, c
        assert "%s **%.2f**" % (names[c], round(V["ratio"][c], 2)) in block2, c
    print("verify: the four flags, eight regions' notified cases either side of their changes and share-notified changes, "
          "eight season-to-date counts and eight week 36 rates recomputed from the bundle; goldens, figures.json and the "
          "write-up carry them")
