"""task309 goldens, written from figures.json: the county workbook retrofit_round_2027.xlsx, the board chart
avoided_loss_by_county.png and the one-page board note board_note_retrofit_2027.docx."""
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

import params as P  # noqa: E402
import postprocess  # noqa: E402

XLSX = "retrofit_round_2027.xlsx"
PNG = "avoided_loss_by_county.png"
DOCX = "board_note_retrofit_2027.docx"
HEAD = PatternFill("solid", fgColor="DCE4EA")
NAMES = [P.COUNTY_NAME[c] for c in P.CANDIDATES]
STORMS = ["Hurricane Corwin", "Tropical Storm Della", "Hurricane Evander"]
STORM_YEAR = {"Hurricane Corwin": 2008, "Tropical Storm Della": 2009, "Hurricane Evander": 2011}


def _hdr(ws, row, cols):
    for j, h in enumerate(cols):
        c = ws.cell(row, j + 1, h)
        c.font = Font(bold=True)
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")


def workbook(path, f):
    wb = Workbook()
    ws = wb.active
    ws.title = "Recommendation"
    ws["A1"] = "2027 FORTIFIED Roof grant round: county comparison"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Catastrophe Risk, for board brief BRD-2026-118. Corvane HU 27.1 event sets, claims to Sep 30, 2026."
    ws["A4"] = "Recommended county"
    ws["B4"] = f["call"]
    ws["B4"].font = Font(bold=True)
    ws["A5"] = "Expected annual loss removed, USD million"
    ws["B5"] = f["call_removed_musd"]
    ws["B5"].number_format = "0.0"
    ws["A6"] = "Runner-up"
    ws["B6"] = f["runner_up"]
    _hdr(ws, 8, ["County", "2027 annual expected loss, standard roofs (USD m)",
                 "Expected annual loss removed by the round (USD m)", "Share of annual loss in events gusting 130 mph+"])
    for i, n in enumerate(NAMES):
        r = 9 + i
        ws.cell(r, 1, n)
        ws.cell(r, 2, f["aal_musd"][n]).number_format = "0.00"
        ws.cell(r, 3, f["removed_musd_2dp"][n]).number_format = "0.00"
        ws.cell(r, 4, f["share_130_up"][n] / 100.0).number_format = "0%"
    ws["A14"] = ("Loss removed: each event's vendor loss times the FORTIFIED Roof cut for the event's peak-gust class, "
                 "from our named-storm claims since 2012 (sheet Cut by gust class), summed at the event's annual rate.")
    ws["A14"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A14:D15")
    for col, wdt in zip("ABCD", (34, 24, 24, 22)):
        ws.column_dimensions[col].width = wdt
    ws.row_dimensions[8].height = 44

    ws = wb.create_sheet("Cut by gust class")
    _hdr(ws, 1, ["Peak gust class (mph)", "FORTIFIED Roof cut in paid loss per dollar of Coverage A"])
    ws.freeze_panes = "A2"
    for i, (k, v) in enumerate(sorted(f["class_cut"].items(), key=lambda t: int(t[0]))):
        ws.cell(2 + i, 1, "%s-%d" % (k, int(k) + 9))
        ws.cell(2 + i, 2, v).number_format = "0.0%"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 34

    ws = wb.create_sheet("MG-04 test")
    _hdr(ws, 1, ["Storm", "Paid on FORTIFIED homes (USD m)", "Class cuts: reproduced vs paid", "One cut for all storms: reproduced vs paid"])
    ws.freeze_panes = "A2"
    for i, sid in enumerate(f["mg04_storms"]):
        ws.cell(2 + i, 1, sid)
        ws.cell(2 + i, 2, f["mg04_paid_musd"][sid]).number_format = "0.0"
        ws.cell(2 + i, 3, round(f["mg04_golden_pct"][sid] / 100.0, 3) + 0.0).number_format = "+0.0%;-0.0%"
        ws.cell(2 + i, 4, round(f["mg04_default_pct"][sid] / 100.0, 3) + 0.0).number_format = "+0.0%;-0.0%"
    ws.cell(8, 1, "MG-04 limit: within 10 percent on every storm with USD 10 million or more paid on FORTIFIED homes.")
    for col, wdt in zip("ABCD", (12, 22, 22, 26)):
        ws.column_dimensions[col].width = wdt

    ws = wb.create_sheet("FORTIFIED homes Sep 30 2026")
    _hdr(ws, 1, ["County", "Insured homes with a valid FORTIFIED Roof certificate", "Coverage A (USD 000)"])
    ws.freeze_panes = "A2"
    for i, n in enumerate(NAMES):
        ws.cell(2 + i, 1, n)
        ws.cell(2 + i, 2, f["ask_a"][n][0]).number_format = "#,##0"
        ws.cell(2 + i, 3, f["ask_a"][n][1]).number_format = "#,##0"
    for col, wdt in zip("ABC", (18, 30, 18)):
        ws.column_dimensions[col].width = wdt

    ws = wb.create_sheet("Storms before 2012")
    _hdr(ws, 1, ["Storm", "Year", "Net paid to date (USD 000)", "Of which paid after Dec 31, 2011 (USD 000)"])
    ws.freeze_panes = "A2"
    for i, s in enumerate(STORMS):
        ws.cell(2 + i, 1, s)
        ws.cell(2 + i, 2, STORM_YEAR[s])
        ws.cell(2 + i, 3, f["ask_b"][s][0]).number_format = "#,##0"
        ws.cell(2 + i, 4, f["ask_b"][s][1]).number_format = "#,##0"
    for col, wdt in zip("ABCD", (24, 8, 24, 30)):
        ws.column_dimensions[col].width = wdt

    ws = wb.create_sheet("Roofs at 15 years")
    _hdr(ws, 1, ["County", "Portfolio homes with a roof 15 years or older on Jun 1, 2027", "Coverage A (USD 000)"])
    ws.freeze_panes = "A2"
    for i, n in enumerate(NAMES):
        ws.cell(2 + i, 1, n)
        ws.cell(2 + i, 2, f["ask_c"][n][0]).number_format = "#,##0"
        ws.cell(2 + i, 3, f["ask_c"][n][1]).number_format = "#,##0"
    for col, wdt in zip("ABC", (18, 34, 18)):
        ws.column_dimensions[col].width = wdt
    wb.save(path)
    postprocess.scrub(path, "SPM-RISK", "2026-11-20")
    postprocess.finish_office(path, "2026-11-12T15:06:00Z", "2026-11-20T21:44:00Z", P.ORG, P.PEOPLE["cat"])


def chart(path, f):
    BLUE, GREY, INK = "#2a6fb0", "#b9c3cc", "#222222"
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9})
    fig, ax = plt.subplots(figsize=(7.2, 3.8), dpi=150)
    xs = range(len(NAMES))
    aal = [f["aal_musd"][n] for n in NAMES]
    rem = [f["removed_musd_2dp"][n] for n in NAMES]
    ax.bar([x - 0.19 for x in xs], aal, width=0.38, color=GREY, label="2027 annual expected loss (standard roofs)")
    cols = [BLUE if n == f["call"] else "#9dbfe0" for n in NAMES]
    ax.bar([x + 0.19 for x in xs], rem, width=0.38, color=cols, label="Expected annual loss the round removes")
    for x, v, n in zip(xs, rem, NAMES):
        ax.text(x + 0.19, v + 0.15, "%.2f" % v, ha="center", va="bottom", fontsize=8.5, color=INK,
                fontweight="bold" if n == f["call"] else "normal")
    ax.set_xticks(list(xs))
    ax.set_xticklabels(NAMES)
    ax.set_ylabel("USD million a year")
    ax.set_title("The round removes the most expected loss in %s, not in the county with the largest annual loss"
                 % f["call"], fontsize=9.5, loc="left")
    fig.text(0.01, 0.01, "Corvane HU 27.1 event losses, gross of reinsurance; FORTIFIED Roof cut by 10 mph gust class "
             "from Sandpiper claims since 2012.", fontsize=7, color="#555555")
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(frameon=False, fontsize=8, loc="upper right")
    fig.tight_layout()
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)


def note(path, f, png):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    d.add_heading("%s should get the 2027 FORTIFIED Roof grant round" % f["call"], level=1)
    d.add_paragraph("To the Board, from Sherry Lang, Catastrophe Risk Manager. November 20, 2026. Re: BRD-2026-118.")
    rem = f["removed_musd_2dp"]
    p = d.add_paragraph()
    p.add_run("Recommendation: award the 2027 round to %s. " % f["call"]).bold = True
    p.add_run("Re-roofing its 1,500 portfolio homes to the FORTIFIED Roof standard removes about USD %.1f million of "
              "expected annual hurricane loss, against USD %.2f million in %s, the runner-up, for the same USD 18.0 "
              "million." % (f["call_removed_musd"], rem[f["runner_up"]], f["runner_up"]))
    d.add_paragraph("%s carries the largest annual expected loss (USD %.2f million) but %d percent of it comes from "
                    "events gusting 130 mph or more at the portfolio. Our claims since 2012 show the FORTIFIED Roof cuts "
                    "paid loss per dollar insured by about %d percent in storms gusting under 130 mph and by about %d "
                    "percent above, where wall, opening and connection failures drive the loss. %s's annual loss sits "
                    "almost entirely below 130 mph." % (
                        f["runner_up"], f["aal_musd"][f["runner_up"]], f["share_130_up"][f["runner_up"]],
                        round(100 * f["cut_under_130"]), round(100 * f["cut_130_up"]), f["call"]))
    d.add_paragraph("The cut by gust class meets MG-04 on all five storms since 2012 with USD 10 million or more paid "
                    "on FORTIFIED homes. A single cut for all storms does not, and ranks the counties by their annual "
                    "loss instead.")
    d.add_picture(png, width=Inches(6.0))
    d.add_paragraph("Detail by county, the MG-04 test table, the FORTIFIED counts at September 30, the pre-2012 storms "
                    "and the portfolio roof ages are in retrofit_round_2027.xlsx.")
    fn = d.add_paragraph("Losses are Corvane HU 27.1's for each county's 1,500 listed homes, gross of reinsurance and "
                         "after the 2 percent named-storm deductible, as the brief compares them; the 2025 storm is "
                         "outside the MG-04 test.")
    for r in fn.runs:
        r.font.size = Pt(8)
    d.core_properties.author = P.PEOPLE["cat"]
    d.core_properties.title = "2027 FORTIFIED Roof grant round"
    d.save(path)
    postprocess.scrub(path, "SPM-RISK", "2026-11-20")
    postprocess.finish_office(path, "2026-11-19T19:12:00Z", "2026-11-20T22:03:00Z", P.ORG, P.PEOPLE["cat"])


def write(outdir, f):
    os.makedirs(outdir, exist_ok=True)
    workbook(os.path.join(outdir, XLSX), f)
    chart(os.path.join(outdir, PNG), f)
    note(os.path.join(outdir, DOCX), f, os.path.join(outdir, PNG))
