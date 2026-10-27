"""Cut the pack: python3 task309/generator/build.py [target dir]. Writes every input file under target/.
Byte-identical on every run."""
import os
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import params as P  # noqa: E402
import postprocess  # noqa: E402
import records as R  # noqa: E402
import world as W  # noqa: E402

TASK = os.path.dirname(HERE)
DISTRACTORS = ["reins", "quotes"]
CONTEXT = ["closeout"]
DATA = ["claims", "legacy", "storms", "footprints", "inforce", "elt", "hazard", "portfolios", "policies", "certs",
        "permits"]
DOCS = ["brief", "mg04", "vendor", "closeout", "reins", "quotes"]
PACK = DATA + DOCS
PRODUCER = "SPM-RISK"


def world():
    w = W.build_world()
    w["by_home"] = {}
    for p in w["policies"]:
        w["by_home"].setdefault(p["home"], []).append(p)
    return w


def footprint_book(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, (s, df) in sheets.items():
        ws = wb.create_sheet(name)
        r0 = 1
        if s["year"] < W.KNOT_UNTIL:
            ws.cell(1, 1, "Gulf Coast Wind Analysis Service").font = Font(bold=True)
            ws.cell(2, 1, "%s (%d): peak 3-second gust at 10 m, open exposure, by ZIP code" % (s["name"], s["year"]))
            ws.cell(3, 1, "Analysis issued %s" % (s["date"] + pd.Timedelta(days=12)).isoformat())
            r0 = 5
        for j, col in enumerate(df.columns):
            ws.cell(r0, j + 1, col).font = Font(bold=True)
        for i, row in enumerate(df.itertuples(index=False)):
            for j, v in enumerate(row):
                ws.cell(r0 + 1 + i, j + 1, v.item() if hasattr(v, "item") else v)
        ws.column_dimensions["A"].width = 11
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["C"].width = 26
    wb.save(path)
    postprocess.scrub(path, PRODUCER, "2025-11-03")
    postprocess.finish_office(path, "2013-03-04T15:20:00Z", "2025-11-03T10:41:00Z", P.ORG, "Sherry Lang")


def build(target):
    os.makedirs(target, exist_ok=True)
    p = lambda k: os.path.join(target, P.FILES[k])
    w = world()
    rng = np.random.default_rng(R.REC_SEED)
    lines = R.claims(w, rng)
    lv, carried = R.legacy(w, rng)
    lines += R.carried_lines(w, carried, rng)
    cl = pd.DataFrame(lines)[["claim_number", "policy_number", "catastrophe_code", "peril", "loss_date", "reported_date",
                              "zip_code", "mitigation_credit", "line_type", "payment_date", "amount_usd",
                              "legacy_claim_ref"]]
    cl = cl.sort_values(["payment_date", "claim_number", "line_type"], kind="mergesort").reset_index(drop=True)
    cl.to_csv(p("claims"), index=False)
    lv.to_csv(p("legacy"), index=False)
    R.storm_register(w).to_csv(p("storms"), index=False)
    footprint_book(p("footprints"), R.footprints(w))
    R.inforce(w).to_parquet(p("inforce"), index=False)
    pf = R.portfolios(w)
    elt, haz = R.vendor(w, pf)
    elt.to_parquet(p("elt"), index=False)
    haz.to_csv(p("hazard"), index=False)
    pf.drop(columns=["_home"]).to_csv(p("portfolios"), index=False)
    R.policy_master(w).to_csv(p("policies"), index=False)
    R.certificates(w, rng).to_csv(p("certs"), index=False)
    R.permits(w, rng).to_csv(p("permits"), index=False)
    try:
        import docs
        docs.build(target, w)
    except ImportError:
        pass
    return w


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(TASK, "target")
    build(out)
