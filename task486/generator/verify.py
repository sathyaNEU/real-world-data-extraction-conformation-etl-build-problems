"""Independent verifier: recomputes the eligibility schedule, every rung of the ladder and both asks from the shipped
pack alone, with its own readers and its own matching, reads back every format the pack ships, and compares the
results with the golden workbook, figures.json and submission.md. It imports nothing from the generator.

    python3 task486/generator/verify.py          exit 0 when every assertion holds
"""
import csv
import datetime as dt
import io
import json
import os
import re
import sys
import zipfile
from collections import defaultdict

import openpyxl
import pandas as pd
from pypdf import PdfReader

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
T = os.path.join(TASK, "target")
GOLD = os.path.join(TASK, "golden")
MARCH = {"22": dt.date(2022, 3, 31), "26": dt.date(2026, 3, 31)}
LISTFILE = {"22": "rating_list_2017_extract_20220404.csv", "26": "rating_list_2023_extract_20260413.csv"}
STAFFED = ("R1", "R2", "R3", "R4", "O1", "W1")
DISTRICTS = ["Rookgate", "Dorrstead", "Glenskarn", "Ingmouth", "Neshdale", "Skelrigg"]
BA = {"5110": "Rookgate", "5120": "Dorrstead", "5130": "Glenskarn", "5140": "Ingmouth", "5150": "Neshdale",
      "5160": "Skelrigg"}
OUT = []


def ok(label, cond, detail=""):
    OUT.append((label, bool(cond), detail))


def path(name):
    return os.path.join(T, name)


def day(v):
    if v in (None, ""):
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%b-%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError("unreadable date %r" % s)


def pdf_text(name):
    return " ".join(" ".join((pg.extract_text() or "").split()) for pg in PdfReader(path(name)).pages)


def bare(name):
    """a ratepayer's or employer's name without its trading style"""
    return re.split(r"\s+t/a\s+", str(name).strip(), flags=re.I)[0].strip()


def canon(name):
    s = bare(name).upper()
    s = re.sub(r"\bLIMITED\b", "LTD", s)
    return " ".join(re.sub(r"[^A-Z0-9&]+", " ", s).split())


def who(name):
    """(initials, SURNAME) of a person or partnership in any of the layouts the files write them"""
    s = bare(name)
    if "," in s:
        sur, given = s.split(",", 1)
        return (given.strip()[:1].upper(), sur.strip().upper())
    w = s.split()
    if "&" in w:
        i = w.index("&")
        return (w[i - 1][0].upper() + "&" + w[i + 1][0].upper(), w[-1].upper())
    return (w[0][0].upper(), w[-1].upper())


# ------------------------------------------------------------------------------------------------ readers

def read_paye():
    a = pd.read_csv(path("paye_returns_m12_2021-22.csv"), dtype=str, keep_default_na=False)
    b = pd.read_csv(path("paye_returns_m12_2025-26.csv"), dtype=str, keep_default_na=False)
    return {"22": pd.DataFrame(dict(name=a["employer_name"], crn=a["company_no"], sic=a["sic2007"],
                                    pc=a["scheme_postcode"], n=a["employees_paid"].astype(int))),
            "26": pd.DataFrame(dict(name=b["employer_name"], crn=b["crn"], sic=b["sic_2007"], pc=b["paye_postcode"],
                                    n=b["paid_employees"].astype(int)))}


def read_lists():
    a = pd.read_csv(path(LISTFILE["22"]), dtype=str, keep_default_na=False)
    b = pd.read_csv(path(LISTFILE["26"]), dtype=str, keep_default_na=False)
    out = {}
    for k, x, cols in (("22", a, ("UARN", "BA_CODE", "BA_REFERENCE", "POSTCODE", "DESCRIPTION_CODE", "TOTAL_AREA",
                                  "RATEABLE_VALUE", "EFFECTIVE_DATE", "EFFECTIVE_TO")),
                       ("26", b, ("uarn", "billing_authority", "ba_ref", "postcode", "primary_desc_code", "floor_area_m2",
                                  "rateable_value", "effective_from", "effective_to"))):
        rows = []
        for r in x.itertuples(index=False):
            r = r._asdict()
            rows.append(dict(uarn=r[cols[0]], district=BA[r[cols[1]]], ref=r[cols[2]], pc=r[cols[3]], d=r[cols[4]],
                             a=float(r[cols[5]]), rv=float(r[cols[6]]), frm=day(r[cols[7]]), to=day(r[cols[8]])))
        out[k] = rows
    return out


def read_accounts(lists):
    """every account on every district's sheet, its property resolved to a UARN by whichever reference the sheet uses"""
    by_ref = {}
    for rows in lists.values():
        for e in rows:
            by_ref[(e["district"], e["ref"])] = e["uarn"]
    wb = openpyxl.load_workbook(path("business_rates_accounts_2017-2026.xlsx"), read_only=True)
    out, sheets = [], {}
    for ws in wb.worksheets:
        if ws.title not in DISTRICTS:
            continue
        it = ws.iter_rows(values_only=True)
        head = [str(h) for h in next(it)]

        def col(*pat):
            hits = [i for i, h in enumerate(head) if any(re.search(p, h, re.I) for p in pat)]
            return hits[0] if hits else None
        c_uarn, c_prop = col(r"^uarn$"), col(r"^prop", r"^property")
        c_name = col(r"liable party", r"ratepayer", r"^name$")
        c_crn = col(r"comp", r"co\. reg")
        c_from, c_to = col(r"start", r"^from$", r"liable_from"), col(r"end", r"^to$", r"liable_to")
        c_occ, c_rel = col(r"^occup"), col(r"relief")
        sheets[ws.title] = head
        for r in it:
            if r[0] is None:
                continue
            u = str(r[c_uarn]) if c_uarn is not None else by_ref[(ws.title, str(r[c_prop]))]
            occ = None if c_occ is None else str(r[c_occ])
            out.append(dict(district=ws.title, uarn=u, name=str(r[c_name]), crn=(r[c_crn] or "") if c_crn is not None else "",
                            frm=day(r[c_from]), to=day(r[c_to]), empty_flag=occ in ("Empty", "N"),
                            relief=r[c_rel] or ""))
    return out, sheets


def read_register():
    x = pd.read_csv(path("companies_register_extract_20260501.csv"), dtype=str, keep_default_na=False)
    reg = {}
    for r in x.itertuples(index=False):
        prev = [re.sub(r"\s*\(TO \d\d/\d\d/\d{4}\)\s*$", "", p).strip() for p in r.previous_names.split(";") if p.strip()]
        reg[r.company_number] = dict(name=r.company_name, prev=prev, sic=set(r.sic_codes.split(";")), psc=r.psc_corporate)
    return reg


def read_schedule():
    x = pd.read_csv(path("town_centre_schedule_2026.csv"), dtype=str, keep_default_na=False)
    centre = dict(zip(x["postcode"], x["centre_code"]))
    names = {}
    for r in x.itertuples(index=False):
        names[r.centre_code] = (r.centre_name, r.district)
    return centre, dict(sorted(names.items())), x


# ------------------------------------------------------------------------------------------------ the analysis

class Island:
    def __init__(self):
        self.paye = read_paye()
        self.lists = read_lists()
        self.accounts, self.sheets = read_accounts(self.lists)
        self.reg = read_register()
        self.centre, self.names, self.sched = read_schedule()
        self.by_name = defaultdict(set)
        for num, c in self.reg.items():
            for nm in [c["name"]] + c["prev"]:
                self.by_name[canon(nm)].add(num)
        self.parent = defaultdict(list)
        for num, c in self.reg.items():
            if c["psc"]:
                self.parent[c["psc"]].append(num)
        self.hist = defaultdict(list)
        for a in self.accounts:
            self.hist[a["uarn"]].append(a)
        for h in self.hist.values():
            h.sort(key=lambda a: a["frm"])
        self.employers = {y: self._employers(y) for y in MARCH}
        self.entries = {y: self._in_force(y) for y in MARCH}
        self.premises = {y: self._premises(y) for y in MARCH}

    def _employers(self, y):
        """each employer's PAYE schemes added up: a company by its number, a sole trader or partnership by initials and
        surname, anyone else by name"""
        out = {}
        for r in self.paye[y].itertuples(index=False):
            if r.crn:
                k = ("C", r.crn)
            elif r.sic.startswith("47"):
                k = ("P",) + who(r.name)
            else:
                k = ("N", canon(r.name))
            e = out.setdefault(k, dict(staff=0, sics=set(), schemes=[]))
            e["staff"] += r.n
            e["sics"].add(r.sic)
            e["schemes"].append((r.pc, r.n))
        return out

    def retail_keys(self, y):
        return [k for k, e in self.employers[y].items() if any(s.startswith("47") for s in e["sics"]) and e["staff"] > 0]

    def _in_force(self, y):
        d = MARCH[y]
        out = {}
        for e in self.lists[y]:
            if e["frm"] <= d and (e["to"] is None or e["to"] >= d):
                assert e["uarn"] not in out, ("two entries in force", e["uarn"])
                out[e["uarn"]] = e
        return out

    def account_on(self, uarn, d):
        cur = [a for a in self.hist.get(uarn, []) if a["frm"] <= d and (a["to"] is None or a["to"] >= d)]
        assert len(cur) <= 1, ("two accounts on one day", uarn, d)
        return cur[0] if cur else None

    def holder(self, a, y):
        """the employer an account is liable for in occupation, or None for an owner's account on an empty premises: one
        on empty property relief, or one in the name of a company the register classes to letting its own real estate
        (68209) or to developing building projects (41100)"""
        if a["empty_flag"] or a["relief"] == "EPR":
            return None
        num = a["crn"]
        if not num:
            hits = self.by_name.get(canon(a["name"]), set())
            assert len(hits) <= 1, ("ambiguous name", a["name"])
            num = next(iter(hits)) if hits else None
        if num is None:
            return ("P",) + who(a["name"])
        c = self.reg.get(num)
        if c and c["sic"] & {"68209", "41100"}:
            for sib in (self.parent.get(c["psc"], []) if c["psc"] else []):
                e = self.employers[y].get(("C", sib))
                if e and any(s.startswith("47") for s in e["sics"]):
                    return ("C", sib)
            return None
        return ("C", num)

    def _premises(self, y):
        out = defaultdict(list)
        self.state = getattr(self, "state", {})
        st = {}
        for uarn, e in self.entries[y].items():
            a = self.account_on(uarn, MARCH[y])
            if a is None:
                st[uarn] = "no account"
                continue
            k = self.holder(a, y)
            st[uarn] = "empty" if k is None else "occupied"
            if k is not None:
                out[k].append(e)
        self.state[y] = st
        return out

    def staff_per_m2(self, y):
        num, den = defaultdict(float), defaultdict(float)
        for k in self.retail_keys(y):
            ps = self.premises[y].get(k, [])
            if len(ps) == 1 and ps[0]["d"] in STAFFED:
                num[ps[0]["d"]] += self.employers[y][k]["staff"]
                den[ps[0]["d"]] += ps[0]["a"]
        return {d: num[d] / den[d] for d in STAFFED}

    def place(self, y, rung):
        """employees placed in each centre under a rung: golden, area, rv, rv23, count, postcode"""
        out = {c: 0.0 for c in self.names}
        dens = self.staff_per_m2(y) if rung == "golden" else None
        rv23 = self.rv_on_2023_list() if rung == "rv23" else None
        for k in self.retail_keys(y):
            e = self.employers[y][k]
            if rung == "postcode":
                for pc, n in e["schemes"]:
                    if pc in self.centre:
                        out[self.centre[pc]] += n
                continue
            ps = self.premises[y].get(k, [])
            if not ps:
                continue
            if rung == "golden":
                w = [p["a"] * dens[p["d"]] for p in ps]
            elif rung == "area":
                w = [p["a"] for p in ps]
            elif rung == "rv":
                w = [p["rv"] for p in ps]
            elif rung == "rv23":
                w = [rv23.get(p["uarn"], p["rv"]) for p in ps]
            else:
                w = [1.0] * len(ps)
            tot = sum(w)
            for p, x in zip(ps, w):
                if p["pc"] in self.centre:
                    out[self.centre[p["pc"]]] += e["staff"] * x / tot
        return out

    def rv_on_2023_list(self):
        cur = {}
        for e in self.lists["26"]:
            if e["to"] is None:
                cur[e["uarn"]] = e["rv"]
        for e in self.lists["26"]:
            cur.setdefault(e["uarn"], e["rv"])
        return cur

    def empty_shops(self, y):
        """per district, the shop premises in its centres empty on the day, and those empty twelve months or more"""
        d = MARCH[y]
        n, longer = defaultdict(int), defaultdict(int)
        for uarn, e in self.entries[y].items():
            if e["d"] != "R1" or e["pc"] not in self.centre:
                continue
            dist = self.names[self.centre[e["pc"]]][1]
            n[dist] += 0
            longer[dist] += 0
            if self.state[y][uarn] != "empty":
                continue
            n[dist] += 1
            h = self.hist[uarn]
            i = next(j for j, a in enumerate(h) if a["frm"] <= d and (a["to"] is None or a["to"] >= d))
            while i > 0 and h[i - 1]["to"] == h[i]["frm"] - dt.timedelta(days=1) and self.holder(h[i - 1], y) is None:
                i -= 1
            if h[i]["frm"] <= dt.date(d.year - 1, 4, 1):
                longer[dist] += 1
        return n, longer


def qualifies(lv22, lv26):
    return {c for c in lv22 if lv26[c] / lv22[c] - 1 <= -0.08}


def r50(x):
    return int(round(x / 50.0)) * 50


# ------------------------------------------------------------------------------------------------ the checks

def formats():
    with open(path("paye_returns_m12_2021-22.csv"), newline="", encoding="utf-8") as fh:
        head = next(csv.reader(fh))
    ok("csv: the 2021-22 PAYE extract is layout 3", head[:3] == ["scheme_ref", "employer_name", "company_no"], head)
    codes = json.load(open(path("valuation_description_codes.json"), encoding="utf-8"))
    cd = {c["code"]: c for c in codes["codes"]}
    ok("json: R1 is shop and premises on net internal area", cd["R1"]["description"].lower() == "shop and premises"
       and cd["R1"]["measurement_basis"] == "NIA", cd.get("R1"))
    lfs = openpyxl.load_workbook(path("labour_force_survey_employment_by_industry_2022_2026.xlsx"), read_only=True)
    cells = [str(v) for ws in lfs.worksheets for r in ws.iter_rows(values_only=True) for v in r if v is not None]
    ok("xlsx: the Labour Force Survey table reads back and is by district, not centre",
       len(cells) > 50 and any("Rookgate" in c for c in cells) and not any(re.fullmatch(r"TC\d\d", c) for c in cells),
       len(cells))
    ff = pd.read_csv(path("town_centre_footfall_daily_2021-2026.csv"), dtype=str)
    ok("csv: the footfall counts read back, every counter in a designated centre", len(ff) > 25000
       and set(ff["centre_code"]) <= {"TC%02d" % i for i in range(1, 32)}, len(ff))
    sic = pd.read_csv(path("sic2007_subclasses.csv"), dtype=str)
    return cd, set(sic["sic2007_subclass"])


def documents():
    pro = pdf_text("town_centre_renewal_fund_2027_prospectus.pdf")
    spec = pdf_text("paye_extract_specification_v4.pdf")
    guide = pdf_text("rating_list_extract_guide.pdf")
    ok("pdf control: the prospectus extracts (its title)", "Town Centre Renewal Fund 2027: prospectus" in pro, len(pro))
    ok("pdf: 3.1 counts employees at the premises where they work",
       "who were paid in the PAYE tax month ending 5 April, counted at the premises where they work" in pro)
    ok("pdf: 3.2 places a premises by its list entry in force on 31 March",
       "the postcode of its entry in the valuation list in force on 31 March of that year is in the centre's schedule" in pro)
    ok("pdf: 3.3 is the 8 per cent test on unrounded figures",
       "8 per cent or more below its workplace retail employment in March 2022, the change being taken on unrounded figures"
       in pro)
    ok("pdf: section 4 rounds employment to the nearest 50 and changes to the whole per cent",
       "rounded to the nearest 50" in pro and "its change to the nearest whole per cent" in pro)
    ok("pdf: the annex defines empty and twelve months or more",
       "no one was in occupation of it on the day" in pro
       and "had been in the list throughout the twelve months to that 31 March and no one was in occupation of it on any "
           "of those days" in pro)
    ok("pdf: the schedule includes terminated postcode units",
       "lists the units the postal administration has terminated since the Order was made" in pro)
    ok("pdf control: the PAYE specification extracts", "PAYE employer extract: specification" in spec, len(spec))
    ok("pdf: a scheme's postcode is a correspondence address, not a workplace",
       "It is not a record of where the employees work." in spec)
    ok("pdf: schemes, not employers, and month 12 runs 6 March to 5 April",
       "the extract lists schemes, not employers" in spec and "Month 12 runs from 6 March to 5 April." in spec)
    ok("pdf control: the list guide extracts", "Rating list extracts: field guide" in guide, len(guide))
    ok("pdf: areas are net internal for shops and gross internal for supermarkets and warehouses",
       "net internal area for shops" in guide and "gross internal area for department stores, supermarkets, retail warehouses"
       in guide)
    return pro


def main():
    cd, sic_titles = formats()
    pro = documents()
    I = Island()
    sched_names = [n for n, _ in I.names.values()]
    ok("pdf: the prospectus annex lists the schedule's 31 centres", len(I.names) == 31
       and all(n in pro for n in sched_names), len(I.names))
    ok("csv: every SIC code in both extracts has a title", set(I.paye["22"]["sic"]) | set(I.paye["26"]["sic"]) <= sic_titles)
    ok("json: every description on both lists is coded", {e["d"] for r in I.lists.values() for e in r} <= set(cd))
    ok("xlsx: six district sheets in six layouts", len(I.sheets) == 6 and len({tuple(h) for h in I.sheets.values()}) == 6)
    acc = openpyxl.load_workbook(path("business_rates_accounts_2017-2026.xlsx"), read_only=True)
    note = " ".join(str(v) for r in acc["Contents"].iter_rows(values_only=True) for v in r if v)
    ok("xlsx: the accounts' contents note files what an account is and what no account means",
       "the occupier's, or the owner's while the property stands empty" in note
       and "has no account while the relief applies" in note, note[-160:])
    ok("xlsx: every account's property resolves to a list entry",
       all(a["uarn"] in {e["uarn"] for r in I.lists.values() for e in r} for a in I.accounts), len(I.accounts))

    mixed = [k for y in MARCH for k, e in I.employers[y].items() if len(e["sics"]) > 1
             and any(x.startswith("47") for x in e["sics"])]
    ok("every retail employer's schemes carry one SIC code", not mixed, mixed[:3])
    # the figures the write-up and the workbook state
    R = {y: I.retail_keys(y) for y in MARCH}
    multi = {y: [k for k in R[y] if len(I.premises[y].get(k, [])) > 1] for y in MARCH}
    lv = {rung: (I.place("22", rung), I.place("26", rung)) for rung in ("golden", "area", "rv", "rv23", "count", "postcode")}
    q = qualifies(*lv["golden"])
    ch = {c: lv["golden"][1][c] / lv["golden"][0][c] - 1 for c in I.names}
    off = [ps for y in MARCH for k in R[y] for ps in I.premises[y].get(k, []) if ps["d"] not in STAFFED]
    ok("every retail employer's premises carry a staffed description", not off, len(off))
    f = json.load(open(os.path.join(HERE, "figures.json")))
    comp = f["components"]
    ok("retail employers in each March", (len(R["22"]), len(R["26"])) == (comp["retail_employers_2022"],
                                                                           comp["retail_employers_2026"]),
       (len(R["22"]), len(R["26"])))
    ok("multi-site retail employers in each March", (len(multi["22"]), len(multi["26"])) ==
       (comp["multi_site_employers_2022"], comp["multi_site_employers_2026"]), (len(multi["22"]), len(multi["26"])))
    me = tuple(sum(I.employers[y][k]["staff"] for k in multi[y]) for y in ("22", "26"))
    ok("multi-site retail employers' employees in each March", me == (comp["multi_site_employees_2022"],
                                                                      comp["multi_site_employees_2026"]), me)
    tot = tuple(int(round(sum(lv["golden"][i].values()) / 1000.0)) for i in (0, 1))
    ok("the 31 centres' workplace retail employment, nearest thousand", tot == (comp["centres_total_2022_thousands"],
                                                                                comp["centres_total_2026_thousands"]), tot)
    names_q = [I.names[c][0] for c in sorted(q)]
    ok("the qualifying centres", sorted(names_q) == sorted(f["call"]["qualifying"]) and len(q) == f["call"]["n"], names_q)
    ok("not on a knife edge: every centre's change at least 0.5 points from the 8 per cent line",
       min(abs(v + 0.08) for v in ch.values()) >= 0.005, round(min(abs(v + 0.08) for v in ch.values()), 4))
    edge = min(min(abs((x % 50) - 25) for x in lv["golden"][i].values()) for i in (0, 1))
    ok("not on a knife edge: every centre's employment at least 1 employee from a rounding boundary", edge >= 1.0,
       round(edge, 2))
    pedge = min(abs((v * 100 % 1) - 0.5) for v in ch.values())
    ok("not on a knife edge: every change at least 0.02 points from a whole-per-cent boundary", pedge >= 0.02, round(pedge, 3))

    # the golden workbook, row by row
    wb = openpyxl.load_workbook(os.path.join(GOLD, "town_centre_fund_2027_eligibility.xlsx"), read_only=True)
    rows = [r for r in wb["Eligibility schedule"].iter_rows(min_row=5, values_only=True) if r[0] and re.fullmatch(r"TC\d\d", str(r[0]))]
    want = [(c, I.names[c][0], I.names[c][1], r50(lv["golden"][0][c]), r50(lv["golden"][1][c]), int(round(ch[c] * 100)),
             "Yes" if c in q else "No") for c in I.names]
    ok("golden workbook: the eligibility schedule, every centre and column", [tuple(r[:7]) for r in rows] == want,
       [(a, b) for a, b in zip(rows, want) if tuple(a[:7]) != b][:3])
    sch = {r["code"]: r for r in f["schedule"]}
    ok("figures.json: the schedule", all((sch[c]["employment_2022"], sch[c]["employment_2026"], sch[c]["change_pct"],
                                          sch[c]["qualifies"]) == w[3:] for c, w in zip(I.names, want)))
    vac = {y: I.empty_shops(y) for y in MARCH}
    ann = {r[0]: r[1:5] for r in wb["Vacancy annex"].iter_rows(min_row=5, values_only=True) if r[0]}
    for dist in DISTRICTS:
        got = (vac["22"][0][dist], vac["22"][1][dist], vac["26"][0][dist], vac["26"][1][dist])
        ok("ask figures, %s: empty and empty twelve months or more, both Marches" % dist,
           tuple(ann[dist]) == got and (f["asks"]["ask_a"][dist]["2022"], f["asks"]["ask_b"][dist]["2022"],
                                        f["asks"]["ask_a"][dist]["2026"], f["asks"]["ask_b"][dist]["2026"]) == got, got)
    ok("golden workbook: the annex's Torvasay row adds up", tuple(ann["Torvasay"]) == tuple(
        sum(ann[d][i] for d in DISTRICTS) for i in range(4)), ann.get("Torvasay"))
    ok("golden workbook: the notes state the rounding", any("nearest 50" in str(v) for r in wb["Notes"].iter_rows(values_only=True)
                                                            for v in r if v))

    # the ladder: each rung files its own wrong set
    for rung in ("area", "rv", "count", "postcode"):
        s = qualifies(*lv[rung])
        L = f["ladder"][rung]
        ok("rung %s: the set figures.json records" % rung, len(s) == L["n"]
           and sorted(I.names[c][0] for c in s - q) == sorted(L["added"])
           and sorted(I.names[c][0] for c in q - s) == sorted(L["dropped"]), (len(s), sorted(s ^ q)))
        ok("rung %s: a wrong answer" % rung, s != q)
    s23 = qualifies(*lv["rv23"])
    ok("rung rv23 (the 2023 list's values for both years): a wrong answer", s23 != q, len(s23))
    ok("the floor-area route is the default path's wrong answer, and wider than the golden", len(qualifies(*lv["area"])) > len(q))

    # what the asks rest on, as the pack shows it
    lone = [e for y in MARCH for u, e in I.entries[y].items() if I.state[y][u] == "no account"]
    ok("every premises with no account on the day is valued under 8,000", lone and max(e["rv"] for e in lone) < 8000,
       (len(lone), max(e["rv"] for e in lone) if lone else None))
    lone_ret = [k for y in MARCH for k in R[y] if any(I.state[y].get(p["uarn"]) != "occupied" for p in I.premises[y].get(k, []))]
    ok("no retail employer's premises is unaccounted or empty", not lone_ret)
    ok("the asks read shops the call does not: no empty shop is a retail employer's", all(
        I.state[y][p["uarn"]] == "occupied" for y in MARCH for k in R[y] for p in I.premises[y].get(k, [])))

    # submission.md block 4 against the recomputation
    sub = open(os.path.join(TASK, "submission.md"), encoding="utf-8").read()
    for c, w in zip(I.names, want):
        line = "%s %s: %s in March 2022, %s in March 2026, %s, %s" % (
            c, w[1], format(w[3], ","), format(w[4], ","), ("%+d per cent" % w[5]) if w[5] else "0 per cent",
            "meets the test" if w[6] == "Yes" else "does not meet the test")
        if line not in sub:
            ok("submission.md: the schedule line for %s" % c, False, line)
            break
    else:
        ok("submission.md: every schedule line", True)
    ok("submission.md: the vacancy lines", all("%s: %d at 31 March 2022, %d of them for twelve months or more; %d at 31 March "
                                               "2026, %d of them" % ((d,) + tuple(ann[d])) in sub for d in DISTRICTS))

    # the cut and the chart
    with zipfile.ZipFile(os.path.join(TASK, "target.zip")) as z:
        same = sorted(z.namelist()) == sorted(os.listdir(T)) and all(
            z.read(n) == open(path(n), "rb").read() for n in z.namelist())
    ok("target.zip holds the pack byte for byte", same)
    png = open(os.path.join(GOLD, "town_centre_retail_employment_2022_2026.png"), "rb").read()
    ok("png: the golden chart is a PNG", png[:8] == b"\x89PNG\r\n\x1a\n" and len(png) > 20000, len(png))
    meta = json.load(open(os.path.join(TASK, "metadata.json")))
    ok("metadata.json: two or more distractors, each shipped", len(meta["distractor_files"]) >= 2
       and all(os.path.exists(os.path.join(TASK, p)) for p in meta["distractor_files"]))

    fails = [o for o in OUT if not o[1]]
    for label, good, detail in OUT:
        print("%s  %s%s" % ("PASS" if good else "FAIL", label, "" if good or detail == "" else "  [%s]" % (detail,)))
    print("%d of %d assertions hold" % (len(OUT) - len(fails), len(OUT)))
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
