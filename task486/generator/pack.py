"""Writes the pack's records from the world and its timeline: the Revenue Office's two month-12 PAYE extracts in their
own layouts, the Valuation Office's 2017 and 2023 list extracts and its description codes, the six districts' rate
accounts in their billing systems' layouts, the companies register, the Department's centre schedule, the SIC 2007
lookup, and the Department's footfall counts and 2025 health check."""
import csv
import datetime as dt
import json
import os

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font

import names as N
import timeline as T
import world as W

BA_CODE = {"RKG": "5110", "DOR": "5120", "GLN": "5130", "ING": "5140", "NSH": "5150", "SKL": "5160"}
SIC_TITLES = {
    "01110": "Growing of cereals (except rice), leguminous crops and oil seeds",
    "10710": "Manufacture of bread; manufacture of fresh pastry goods and cakes",
    "10890": "Manufacture of other food products n.e.c.", "16100": "Sawmilling and planing of wood",
    "25110": "Manufacture of metal structures and parts of structures", "25620": "Machining",
    "28990": "Manufacture of other special-purpose machinery n.e.c.", "33120": "Repair of machinery",
    "41201": "Construction of commercial buildings", "41202": "Construction of domestic buildings",
    "43210": "Electrical installation", "43220": "Plumbing, heat and air-conditioning installation",
    "43320": "Joinery installation", "43341": "Painting", "43390": "Other building completion and finishing",
    "45200": "Maintenance and repair of motor vehicles", "46390": "Non-specialised wholesale of food, beverages and tobacco",
    "46710": "Wholesale of solid, liquid and gaseous fuels and related products", "46900": "Non-specialised wholesale trade",
    "47110": "Retail sale in non-specialised stores with food, beverages or tobacco predominating",
    "47190": "Other retail sale in non-specialised stores",
    "47210": "Retail sale of fruit and vegetables in specialised stores",
    "47220": "Retail sale of meat and meat products in specialised stores",
    "47230": "Retail sale of fish, crustaceans and molluscs in specialised stores",
    "47240": "Retail sale of bread, cakes, flour confectionery and sugar confectionery in specialised stores",
    "47250": "Retail sale of beverages in specialised stores", "47260": "Retail sale of tobacco products in specialised stores",
    "47290": "Other retail sale of food in specialised stores",
    "47410": "Retail sale of computers, peripheral units and software in specialised stores",
    "47421": "Retail sale of mobile telephones", "47510": "Retail sale of textiles in specialised stores",
    "47520": "Retail sale of hardware, paints and glass in specialised stores",
    "47530": "Retail sale of carpets, rugs, wall and floor coverings in specialised stores",
    "47540": "Retail sale of electrical household appliances in specialised stores",
    "47591": "Retail sale of musical instruments and scores",
    "47599": "Retail of furniture, lighting, and similar (not musical instruments or scores) in specialised store",
    "47610": "Retail sale of books in specialised stores", "47620": "Retail sale of newspapers and stationery in specialised stores",
    "47640": "Retail sale of sports goods, fishing gear, camping goods, boats and bicycles",
    "47650": "Retail sale of games and toys in specialised stores", "47710": "Retail sale of clothing in specialised stores",
    "47721": "Retail sale of footwear in specialised stores", "47730": "Dispensing chemist in specialised stores",
    "47741": "Retail sale of hearing aids",
    "47749": "Retail sale of medical and orthopaedic goods in specialised stores (not incl. hearing aids) n.e.c.",
    "47750": "Retail sale of cosmetic and toilet articles in specialised stores",
    "47760": "Retail sale of flowers, plants, seeds, fertilisers, pet animals and pet food in specialised stores",
    "47770": "Retail sale of watches and jewellery in specialised stores",
    "47781": "Retail sale in commercial art galleries", "47782": "Retail sale by opticians",
    "47789": "Other retail sale of new goods in specialised stores (not commercial art galleries and opticians)",
    "47791": "Retail sale of antiques including antique books in stores",
    "47799": "Retail sale of other second-hand goods in stores (not incl. antiques)",
    "47910": "Retail sale via mail order houses or via Internet", "49320": "Taxi operation",
    "49410": "Freight transport by road", "52103": "Operation of warehousing and storage facilities for land transport activities",
    "55100": "Hotels and similar accommodation", "56101": "Licensed restaurants", "56102": "Unlicensed restaurants and cafes",
    "56103": "Take-away food shops and mobile food stands", "56302": "Public houses and bars",
    "62012": "Business and domestic software development", "62020": "Information technology consultancy activities",
    "64191": "Banks", "64192": "Building societies", "66220": "Activities of insurance agents and brokers",
    "41100": "Development of building projects",
    "68209": "Other letting and operating of own or leased real estate", "68310": "Real estate agencies",
    "69102": "Solicitors", "69201": "Accounting and auditing activities", "69202": "Bookkeeping activities",
    "70100": "Activities of head offices", "70229": "Management consultancy activities other than financial management",
    "71111": "Architectural activities", "71129": "Other engineering activities", "74100": "Specialised design activities",
    "78200": "Temporary employment agency activities", "80100": "Private security activities",
    "81210": "General cleaning of buildings", "82990": "Other business support service activities n.e.c.",
    "84110": "General public administration activities", "85200": "Primary education",
    "85310": "General secondary education", "86101": "Hospital activities", "86210": "General medical practice activities",
    "86230": "Dental practice activities", "87300": "Residential care activities for the elderly and disabled",
    "88100": "Social work activities without accommodation for the elderly and disabled",
    "88990": "Other social work activities without accommodation n.e.c.", "92000": "Gambling and betting activities",
    "93110": "Operation of sports facilities", "93130": "Fitness facilities", "94910": "Activities of religious organisations",
    "96020": "Hairdressing and other beauty treatment", "96090": "Other personal service activities n.e.c.",
}


def _d(x, fmt):
    if x is None:
        return ""
    return x.strftime(fmt)


def rv_round(v):
    if v < 10000:
        return int(round(v / 25.0)) * 25
    if v < 100000:
        return int(round(v / 50.0)) * 50
    return int(round(v / 250.0)) * 250


def ba_ref(district, n):
    return dict(RKG="RG%07d" % n, DOR="D/%08d" % n, GLN="GLN-%06d" % n, ING="I%07d" % n, NSH="N/%06d" % n,
                SKL="SK%07d" % n)[district]


# ------------------------------------------------------------------------------------------------ PAYE

def _paye_name(E, year):
    nm = E["name"]
    if E.get("prev") and year == "22":
        old, when = E["prev"][0]
        if when > dt.date(2022, 5, 16):
            nm = old
    return nm


def paye(w, path22, path26):
    rows22, rows26 = [], []
    for sc in w.schemes:
        E = w.E[sc["eid"]]
        sic = E.get("sic", "")
        crn = E.get("company_no", "") if E.get("incorporated") else ""
        if sc["in22"]:
            rows22.append(["%s/%s" % (sc["office"], sc["ref22"]), _paye_name(E, "22").upper(), crn, sic,
                           E["held"]["22"][0], sc["s22"], "2021-22 M12"])
        if sc["in26"]:
            trading = E.get("brand_name") or E.get("trading") or ""
            rows26.append([sc["office"], sc["ref26"], E["name"], trading if trading != E["name"] else "", crn, sic,
                           E["held"]["26"][0], sc["s26"], "2025/26 M12"])
    rows22.sort(key=lambda r: r[0])
    rows26.sort(key=lambda r: (r[0], r[1]))
    with open(path22, "w", newline="", encoding="utf-8") as fh:
        c = csv.writer(fh)
        c.writerow(["scheme_ref", "employer_name", "company_no", "sic2007", "scheme_postcode", "employees_paid", "return_period"])
        c.writerows(rows22)
    with open(path26, "w", newline="", encoding="utf-8") as fh:
        c = csv.writer(fh)
        c.writerow(["tax_office", "employer_ref", "employer_name", "trading_name", "crn", "sic_2007", "paye_postcode",
                    "paid_employees", "period"])
        c.writerows(rows26)
    return len(rows22), len(rows26)


# ------------------------------------------------------------------------------------------------ lists

def lists(w, path17, path23, pathcodes):
    with open(path17, "w", newline="", encoding="utf-8") as fh:
        c = csv.writer(fh)
        c.writerow(["UARN", "BA_CODE", "BA_REFERENCE", "NUMBER_OR_NAME", "STREET", "TOWN", "POSTCODE", "DESCRIPTION_CODE",
                    "DESCRIPTION", "TOTAL_AREA", "AREA_BASIS", "RATEABLE_VALUE", "EFFECTIVE_DATE", "EFFECTIVE_TO"])
        for e in sorted(w.entries17, key=lambda e: (BA_CODE[e["district"]], e["uarn"], e["frm"])):
            c.writerow([e["uarn"], BA_CODE[e["district"]], ba_ref(e["district"], e["ba_n"]), e["number"], e["street"].upper(),
                        e["locality"].upper(), e["pc"], e["d"], W.DESC[e["d"]][0].upper(), "%.2f" % e["a"], W.DESC[e["d"]][1],
                        rv_round(e["rv"]), _d(e["frm"], "%d/%m/%Y"), _d(e["to"], "%d/%m/%Y")])
    with open(path23, "w", newline="", encoding="utf-8") as fh:
        c = csv.writer(fh)
        c.writerow(["uarn", "billing_authority", "ba_ref", "address", "postcode", "primary_desc_code", "primary_description",
                    "floor_area_m2", "measurement_basis", "rateable_value", "effective_from", "effective_to"])
        for e in sorted(w.entries23, key=lambda e: (BA_CODE[e["district"]], e["uarn"], e["frm"])):
            c.writerow([e["uarn"], BA_CODE[e["district"]], ba_ref(e["district"], e["ba_n"]),
                        "%s %s, %s" % (e["number"], e["street"], e["locality"]), e["pc"], e["d"], W.DESC[e["d"]][0],
                        "%.1f" % e["a"], W.DESC[e["d"]][1], rv_round(e["rv"]), _d(e["frm"], "%Y-%m-%d"), _d(e["to"], "%Y-%m-%d")])
    codes = dict(publisher="Torvasay Valuation Office", dataset="Primary description codes, rating lists 2017 and 2023",
                 updated="2023-02-14",
                 codes=[dict(code=k, description=v[0], measurement_basis=v[1]) for k, v in sorted(W.DESC.items())])
    json.dump(codes, open(pathcodes, "w", encoding="utf-8"), indent=1)
    return len(w.entries17), len(w.entries23)


# ------------------------------------------------------------------------------------------------ accounts

LAYOUT = {
    "RKG": dict(sheet="Rookgate", cols=["Account Number", "UARN", "Liable Party", "Company Number", "Liability Start",
                                         "Liability End", "Relief"], date="%d/%m/%Y", ref="uarn", crn=True, status=False),
    "GLN": dict(sheet="Glenskarn", cols=["Acct", "Property Reference", "Ratepayer", "Co. Reg. No.", "From", "To",
                                         "Reliefs"], date="%d/%m/%Y", ref="ba", crn=True, status=False),
    "DOR": dict(sheet="Dorrstead", cols=["Account No", "Prop Ref", "Ratepayer Name", "Start Date", "End Date", "Relief Code"],
                date="%d/%m/%Y", ref="ba", crn=False, status=False),
    "ING": dict(sheet="Ingmouth", cols=["account_ref", "uarn", "ratepayer", "company_no", "liable_from", "liable_to", "relief"],
                date="%d-%b-%Y", ref="uarn", crn=True, status=False),
    "NSH": dict(sheet="Neshdale", cols=["Account", "Property", "Name", "Liability Start", "Liability End", "Relief Type"],
                date="%d/%m/%Y", ref="ba", crn=False, status=False),
    "SKL": dict(sheet="Skelrigg", cols=["AccountId", "UARN", "RatepayerName", "CompanyRegNo", "StartDate", "EndDate", "ReliefCode"],
                date="%Y-%m-%d", ref="uarn", crn=True, status=False),
}


def styled(district, name, person=None, partnership=False, trading=None):
    """a ratepayer's name as each district's billing system writes it"""
    def ltd(x):
        return x.replace(" Limited", " Ltd")
    if person is not None:
        f, s = person
        if district == "RKG":
            out = name
        elif district == "GLN":
            out = name if partnership else "%s %s" % (f[0], s)
        elif district == "DOR":
            out = (name if partnership else "%s %s" % (f[0], s)).upper()
        elif district == "ING":
            out = name if partnership else "%s, %s" % (s, f)
        elif district == "NSH":
            out = name if partnership else "%s %s" % (f[0], s)
        else:
            out = name.upper()
    else:
        out = dict(RKG=name, GLN=ltd(name), DOR=ltd(name).upper(), ING=name, NSH=ltd(name), SKL=name.upper())[district]
    if trading and district in ("DOR", "NSH") and trading.lower() not in out.lower():
        out += (" T/A " if district == "DOR" else " t/a ") + (trading.upper() if district == "DOR" else trading)
    return out


def _charged(lo, hi, entries):
    """the parts of a trader's spell from lo to hi (None: open) on which the list entry in force values the premises
    at or over the relief's ceiling, an entry's value standing from the spell's start until it takes effect"""
    if not entries or (hi is not None and lo > hi):
        return []
    entries = sorted(entries, key=lambda e: e["frm"])
    out = []
    for i, e in enumerate(entries):
        a = lo if i == 0 else max(lo, e["frm"])
        b = entries[i + 1]["frm"] - T.DAY if i + 1 < len(entries) else hi
        if hi is not None and b is not None:
            b = min(b, hi)
        if (b is not None and a > b) or rv_round(e["rv"]) < T.STR_CEILING:
            continue
        if out and out[-1][1] is not None and out[-1][1] + T.DAY == a:
            out[-1] = (out[-1][0], b)
        else:
            out.append((a, b))
    return out


def accounts(w, path, rng):
    """one row per account with a liability: an occupier's, or an owner's while the premises stood empty"""
    wb = openpyxl.Workbook()
    contents = wb.active
    contents.title = "Contents"
    rows = {d: [] for d in BA_CODE}
    counter = {d: int(rng.integers(100000, 400000)) for d in BA_CODE}
    landlord_by = {l["lid"]: l for l in w.landlords}
    prior = {}
    w.account_truth = []
    allp = list(w.P) + list(w.deleted)
    traders = {}
    ents = {"17": {}, "23": {}}
    for k, src in (("17", w.entries17), ("23", w.entries23)):
        for e in src:
            ents[k].setdefault(e["uarn"], []).append(e)
    for p in allp:
        d = p["district"]
        lay = LAYOUT[d]
        ref = str(p["uarn"]) if lay["ref"] == "uarn" else ba_ref(d, p["ba_n"])
        for sp in p["spells"]:
            kind, who = sp["who"]
            frm, to = sp["frm"], sp["to"]
            if to is not None and to < T.D0:
                continue
            pieces = [(frm, to)]
            if kind == "t":
                # a trader with no employees pays nothing while the list values its premises below the relief's
                # ceiling, so it has no account then
                pieces = (_charged(frm, min(to or T.END, T.L23 - T.DAY), ents["17"].get(p["uarn"], []))
                          + _charged(max(frm, T.L23), to, ents["23"].get(p["uarn"], [])))
            for a, b in pieces:
                relief, status = "", "Occupied"
                company_no = ""
                if kind == "e":
                    E = w.E[who]
                    if E["kind"] == "chain" and who in w.groupprop and p["d"] == "R3":
                        name = "Kestle Properties Limited"
                        company_no = w.groupprop[who]
                        person = None
                        trading = "Kestle"
                    else:
                        name = E["name"]
                        if E.get("prev") and a < E["prev"][0][1]:
                            name = E["prev"][0][0]
                        company_no = E.get("company_no", "") if E.get("incorporated") else ""
                        person = E.get("person")
                        trading = E.get("brand_name") or E.get("trading")
                    rvn = p.get("rv23") if (b is None or b >= T.L23) else p.get("rv17")
                    if E.get("sic") == "88990":
                        relief = "CHR"
                    elif E["kind"] in ("ind", "ind_out") and rvn is not None and rv_round(rvn) < 15000:
                        relief = "SBR"
                    styled_name = styled(d, name, person, E.get("partnership"), trading if rng.random() < 0.7 else None)
                elif kind == "t":
                    if who not in traders:
                        f, s = w.namer.person(avoid=True)
                        traders[who] = (f, s, "%s's %s" % (s, ["Crafts", "Gifts", "Repairs", "Hair", "Tailoring", "Prints"][int(rng.integers(0, 6))]))
                    f, s, trading = traders[who]
                    styled_name = styled(d, "%s %s" % (f, s), (f, s), False, trading)
                    rvn = p.get("rv23") if (b is None or b >= T.L23) else p.get("rv17")
                    if rvn is not None and rv_round(rvn) < 15000:
                        relief = "SBR"
                elif kind == "x":
                    L = landlord_by[sp["owner"]]
                    styled_name = styled(d, L["name"])
                    company_no = L["company_no"]
                    status = "Empty"
                    months = 6 if p["d"] in T.INDUSTRIAL else 3
                    span_end = b or T.END
                    if (span_end - a).days < months * 30.4 and not sp.get("sold"):
                        relief = "EPR"
                else:
                    key = p["uarn"]
                    if key not in prior:
                        if rng.random() < 0.5:
                            f, s = w.namer.person(avoid=True)
                            prior[key] = (styled(d, "%s %s" % (f, s), (f, s)), "")
                        else:
                            f, s = w.namer.person(avoid=True)
                            nm = w.namer.company("%s %s Limited" % (s, ["Trading", "Retail", "Services", "Enterprises"][int(rng.integers(0, 4))]))
                            num = w.company(nm, "47190", status="Dissolved", dissolved=a + dt.timedelta(days=int(rng.integers(400, 2400))))
                            prior[key] = (styled(d, nm), num)
                    styled_name, company_no = prior[key]
                counter[d] += int(rng.integers(1, 40))
                acct = dict(RKG="10%06d" % counter[d], GLN="%08d" % (counter[d] * 3), DOR="D%07d" % counter[d],
                            ING="ING/%06d" % counter[d], NSH="%07d" % counter[d], SKL="SK-%06d" % counter[d])[d]
                fmt = lay["date"]
                rowv = [acct, ref, styled_name]
                if lay["crn"]:
                    rowv.append(company_no)
                rowv += [_d(a, fmt), _d(b, fmt)]
                if lay["status"]:
                    rowv.append(status if d == "RKG" else ("Y" if status == "Occupied" else "N"))
                rowv.append(relief)
                rows[d].append((ref, a, rowv))
                w.account_truth.append(dict(pid=p.get("pid"), uarn=p["uarn"], district=d, frm=a, to=b, kind=kind,
                                            who=who, relief=relief))
    for d in ("RKG", "DOR", "GLN", "ING", "NSH", "SKL"):
        lay = LAYOUT[d]
        ws = wb.create_sheet(lay["sheet"])
        ws.append(lay["cols"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for _, _, r in sorted(rows[d], key=lambda x: (x[0], x[1])):
            ws.append([None if v == "" else v for v in r])
        ws.freeze_panes = "A2"
        for i, col in enumerate(lay["cols"]):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = max(12, min(42, len(col) + 6)) if i != 2 else 44
    contents.append(["Business rates accounts supplied to Torvasay Statistics under the 2024 data-sharing arrangement"])
    contents.append([])
    contents.append(["Sheet", "Billing authority", "System", "Extract run", "Accounts from"])
    systems = dict(RKG="Civica Rates 7", DOR="NDR Plus", GLN="Civica Rates 7", ING="Revenue Manager", NSH="NDR Plus", SKL="RatesOnline")
    for d in ("RKG", "DOR", "GLN", "ING", "NSH", "SKL"):
        contents.append([LAYOUT[d]["sheet"], "%s District Council" % dict(N.DISTRICTS_NAMES)[d], systems[d],
                         "0%d/05/2026" % (4 + list(BA_CODE).index(d) % 3), "01/04/2017"])
    contents.append([])
    contents.append(["Each sheet lists the accounts with a liability on a property: the occupier's, or the owner's while the "
                     "property stands empty."])
    contents.append(["An occupier whose bill is wholly relieved under small trader relief (rateable value under 8,000) has "
                     "no account while the relief applies."])
    contents["A1"].font = Font(bold=True)
    for c in contents[3]:
        c.font = Font(bold=True)
    contents.column_dimensions["A"].width = 14
    contents.column_dimensions["B"].width = 30
    contents.column_dimensions["C"].width = 18
    contents.column_dimensions["D"].width = 13
    contents.column_dimensions["E"].width = 14
    wb.save(path)
    return {d: len(v) for d, v in rows.items()}


# ------------------------------------------------------------------------------------------------ companies

def companies(w, path, rng):
    reg_pc = {}
    for e, E in w.E.items():
        if E.get("incorporated") and E.get("company_no"):
            held = E.get("held", {}).get("26") or E.get("held", {}).get("22")
            reg_pc[E["company_no"]] = held[0] if held else w.pc_unit("RKG")
    with open(path, "w", newline="", encoding="utf-8") as fh:
        c = csv.writer(fh)
        c.writerow(["company_number", "company_name", "previous_names", "company_status", "incorporation_date",
                    "dissolution_date", "registered_office_postcode", "sic_codes", "psc_corporate"])
        for co in sorted(w.companies, key=lambda x: x["number"]):
            prev = "; ".join("%s (to %s)" % (n, (d - T.DAY).strftime("%d/%m/%Y")) for n, d in co["prev"])
            pc = reg_pc.get(co["number"]) or w.pc_unit(["RKG", "DOR", "GLN", "ING", "NSH", "SKL"][int(rng.integers(0, 6))])
            c.writerow([co["number"], co["name"].upper(), prev.upper(), co["status"], co["incorporated"].strftime("%d/%m/%Y"),
                        _d(co.get("dissolved"), "%d/%m/%Y"), pc, co["sic"], co.get("psc") or ""])
    return len(w.companies)


# ------------------------------------------------------------------------------------------------ the schedule, SIC

def schedule(w, path):
    rows = []
    for c in w.centres:
        for pc, st in w.centre_units[c["k"]]:
            if pc in w.recoded:
                rows.append([c["code"], c["name"], dict(N.DISTRICTS_NAMES)[c["district"]], pc, "terminated", "31/03/2023"])
                rows.append([c["code"], c["name"], dict(N.DISTRICTS_NAMES)[c["district"]], w.recoded[pc], "live", ""])
            else:
                rows.append([c["code"], c["name"], dict(N.DISTRICTS_NAMES)[c["district"]], pc, "live", ""])
    rows.sort(key=lambda r: (r[0], r[3]))
    with open(path, "w", newline="", encoding="utf-8") as fh:
        cw = csv.writer(fh)
        cw.writerow(["centre_code", "centre_name", "district", "postcode", "postcode_status", "terminated_on"])
        cw.writerows(rows)
    return len(rows)


def sic_lookup(w, path):
    used = set(E.get("sic") for E in w.E.values() if E.get("sic")) | set(co["sic"] for co in w.companies)
    missing = sorted(u for u in used if u not in SIC_TITLES)
    assert not missing, missing
    with open(path, "w", newline="", encoding="utf-8") as fh:
        c = csv.writer(fh)
        c.writerow(["sic2007_subclass", "title"])
        for k in sorted(used):
            c.writerow([k, SIC_TITLES[k]])
    return len(used)
