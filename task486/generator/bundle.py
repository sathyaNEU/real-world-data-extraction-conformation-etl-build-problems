"""The cut pack read as a solver reads it, and every construction on it: the golden (the unit's findings resolved into
people strongest first, refusing any link that would put two records of one source in one person), the default path
(every chain of links merged), the rungs between and around them, the 2027 split, and the two asks. The golden
workbook, the checks and the pair ceiling all compute from here; verify.py recomputes the graded figures with its own
reader."""
import collections
import csv
import datetime as dt
import json
import os

import openpyxl

LO, HI = dt.date(2025, 10, 1), dt.date(2026, 9, 30)
QUARTERS = [(dt.date(2025, 10, 1), dt.date(2025, 12, 31)), (dt.date(2026, 1, 1), dt.date(2026, 3, 31)),
            (dt.date(2026, 4, 1), dt.date(2026, 6, 30)), (dt.date(2026, 7, 1), dt.date(2026, 9, 30))]
COCS = ["HAL", "MEL", "OND", "PYR"]
RECIPIENTS = COCS + ["DVC"]
FUND, ADMIN = 18_750_000, 0.04
HOMELESS = ("Street", "Homeless Shelter", "Transitional")
TYPES = (True, False)                   # households with children, households of adults only
TYPE_NAME = {True: "Households with children", False: "Adults without children"}
FILES = {"client": "hmis_client_fy2025_fy2026.csv", "enrollment": "hmis_enrollment_fy2025_fy2026.csv",
         "exit": "hmis_exit_fy2025_fy2026.csv", "services": "hmis_services_fy2025_fy2026.csv",
         "dvc": "dvc_comparable_db_export.xlsx", "vsm": "wsh_hch_encounters_fy2025_fy2026.csv",
         "hot": "hot_field_app_contacts.jsonl", "pairs": "sdlu_linkage_pairs_run_20261008.csv",
         "register": "partnership_project_register.xlsx"}


def hdate(s):
    """An HMIS date as the vendors export it: ISO, or Melbray's month/day/year."""
    if not s:
        return None
    if "/" in s:
        m, d, y = (int(v) for v in s.split("/"))
        return dt.date(y, m, d)
    return dt.date.fromisoformat(s[:10])


def _csv(path):
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


class Inputs:
    def __init__(self, target):
        p = lambda k: os.path.join(target, FILES[k])
        wb = openpyxl.load_workbook(p("register"), read_only=True)
        self.ptype = {}
        for r in list(wb["HMIS projects"].iter_rows(values_only=True))[1:]:
            self.ptype[int(r[1])] = int(r[3])
        self.site_coc = {r[0]: r[2] for r in list(wb["Clinic sites"].iter_rows(values_only=True))[1:]}
        self.clients = [("HMIS-" + r["SourceCoC"], r["PersonalID"]) for r in _csv(p("client"))]
        exits = {(r["SourceCoC"], r["EnrollmentID"]): hdate(r["ExitDate"]) for r in _csv(p("exit"))}
        self.enrollments = []
        for r in _csv(p("enrollment")):
            self.enrollments.append({
                "coc": r["SourceCoC"], "eid": r["EnrollmentID"], "pid": r["PersonalID"], "project": int(r["ProjectID"]),
                "type": self.ptype[int(r["ProjectID"])], "entry": hdate(r["EntryDate"]),
                "exit": exits.get((r["SourceCoC"], r["EnrollmentID"])), "hh": r["HouseholdID"],
                "rel": int(r["RelationshipToHoH"]), "move_in": hdate(r["MoveInDate"])})
        self.beds = [{"coc": r["SourceCoC"], "pid": r["PersonalID"], "eid": r["EnrollmentID"],
                      "date": hdate(r["DateProvided"]), "rt": int(r["RecordType"])} for r in _csv(p("services"))]
        wb = openpyxl.load_workbook(p("dvc"), read_only=True)
        self.dv_clients = [r[0] for r in list(wb["Clients"].iter_rows(values_only=True))[1:]]
        self.dv_stays = []
        for r in list(wb["Stays"].iter_rows(values_only=True))[1:]:
            self.dv_stays.append({"key": r[1], "start": r[3].date(), "end": r[4].date() if r[4] else None})
        self.encounters = []
        for r in _csv(p("vsm")):
            m, d, y = (int(v) for v in r["EncounterDate"].split("/"))
            self.encounters.append({"key": r["PatientKey"], "date": dt.date(y, m, d), "site": r["Site"],
                                    "status": r["HousingStatus"]})
        self.contacts = []
        with open(p("hot"), encoding="utf-8") as fh:
            for line in fh:
                c = json.loads(line)
                self.contacts.append({"key": c["client_key"], "date": dt.date.fromisoformat(c["contact_time"][:10]),
                                      "time": c["contact_time"]})
        self.pairs = []
        for r in _csv(p("pairs")):
            self.pairs.append(((r["source_a"], r["record_a"]), (r["source_b"], r["record_b"]), float(r["score"]),
                               r["finding"]))
        recs = list(self.clients) + [("DVC", k) for k in self.dv_clients] + \
            [("WSH", k) for k in sorted({e["key"] for e in self.encounters})] + \
            [("HOT", k) for k in sorted({c["key"] for c in self.contacts})]
        self.records = recs
        self.index = {r: i for i, r in enumerate(recs)}
        self.links = [(self.index[a], self.index[b], s) for a, b, s, f in self.pairs if f == "LINK"]
        self.nonlinks = [(self.index[a], self.index[b], s) for a, b, s, f in self.pairs if f == "NONLINK"]

    # -------------------------------------------------------------------------------------------- activity
    def activity(self, dv_end_inclusive=False, clinic_any_status=False, hot_utc=False, exit_night=False):
        """Per record: whether it shows FY2026 service under the policy, and what placement reads from it. The
        keywords are the alternative readings the fork table asserts the data makes agree."""
        act = [dict(active=False, hmis_entry=None, hmis_coc=None, dv=False, last_other=None, other_coc=None)
               for _ in self.records]
        for e in self.enrollments:
            i = self.index[("HMIS-" + e["coc"], e["pid"])]
            open_ = e["entry"] <= HI and (e["exit"] is None or (e["exit"] > LO if exit_night else e["exit"] >= LO))
            if open_:
                a = act[i]
                a["active"] = True
                if a["hmis_entry"] is None or e["entry"] > a["hmis_entry"]:
                    a["hmis_entry"], a["hmis_coc"] = e["entry"], e["coc"]
        for s in self.dv_stays:
            i = self.index[("DVC", s["key"])]
            end_ok = s["end"] is None or (s["end"] >= LO if dv_end_inclusive else s["end"] > LO)
            if s["start"] <= HI and end_ok:
                act[i]["active"] = True
                act[i]["dv"] = True
        for e in self.encounters:
            if LO <= e["date"] <= HI:
                i = self.index[("WSH", e["key"])]
                a = act[i]
                if clinic_any_status or e["status"] in HOMELESS:
                    a["active"] = True
                if a["last_other"] is None or e["date"] > a["last_other"]:
                    a["last_other"], a["other_coc"] = e["date"], self.site_coc[e["site"]]
        for c in self.contacts:
            d = c["date"]
            if hot_utc:
                t = dt.datetime.fromisoformat(c["time"])
                d = t.astimezone(dt.timezone.utc).date()
            if LO <= d <= HI:
                i = self.index[("HOT", c["key"])]
                a = act[i]
                a["active"] = True
                if a["last_other"] is None or d > a["last_other"]:
                    a["last_other"], a["other_coc"] = d, "HAL"
        return act

    # -------------------------------------------------------------------------------------------- people
    def resolve(self, constrained=True, links=None, order=None, keep=None):
        """Root per record: the unit's LINK findings taken strongest first (ties by record order unless an order is
        given); constrained refuses a link that would put two records of one source in one person. keep limits the
        records whose links are used (the FY2026-only rung)."""
        n = len(self.records)
        par = list(range(n))
        srcs = [{self.records[i][0]} for i in range(n)]

        def f(x):
            while par[x] != x:
                par[x] = par[par[x]]
                x = par[x]
            return x
        seq = order if order is not None else sorted(links if links is not None else self.links,
                                                     key=lambda t: (-t[2], t[0], t[1]))
        for i, j, s in seq:
            if keep is not None and (i not in keep or j not in keep):
                continue
            a, b = f(i), f(j)
            if a == b:
                continue
            if constrained and (srcs[a] & srcs[b]):
                continue
            par[a] = b
            srcs[b] |= srcs[a]
        return [f(i) for i in range(n)]

    def quarantine(self):
        """Over-cleaning: a record linked to two records of one source loses every link, and the rest are merged."""
        nb = collections.defaultdict(list)
        for i, j, s in self.links:
            nb[i].append(j)
            nb[j].append(i)
        bad = {i for i, js in nb.items() if any(v > 1 for v in collections.Counter(self.records[j][0] for j in js).values())}
        kept = [l for l in self.links if l[0] not in bad and l[1] not in bad]
        return self.resolve(constrained=False, links=kept)

    def served(self, roots, act):
        groups = collections.defaultdict(list)
        for i, r in enumerate(roots):
            groups[r].append(i)
        out = {}
        for root, idx in groups.items():
            if not any(act[i]["active"] for i in idx):
                continue
            hm = [(act[i]["hmis_entry"], act[i]["hmis_coc"]) for i in idx if act[i]["hmis_entry"] is not None]
            if hm:
                out[root] = max(hm)[1]
            elif any(act[i]["dv"] for i in idx):
                out[root] = "DVC"
            else:
                ot = [(act[i]["last_other"], act[i]["other_coc"]) for i in idx if act[i]["last_other"] is not None]
                out[root] = max(ot)[1]
        return out

    @staticmethod
    def counts(served):
        c = collections.Counter(served.values())
        return {k: c.get(k, 0) for k in RECIPIENTS}


def allocation(cnt):
    pool = FUND * (1 - ADMIN)
    total = sum(cnt.values())
    return {k: int(round(pool * cnt[k] / total)) for k in RECIPIENTS}


def allocation_exact(cnt):
    pool = FUND * (1 - ADMIN)
    total = sum(cnt.values())
    return {k: pool * cnt[k] / total for k in RECIPIENTS}


# ------------------------------------------------------------------------------------------------ the asks
def _q(d):
    for k, (a, b) in enumerate(QUARTERS):
        if a <= d <= b:
            return k
    return None


def shelter_nights(I, dedup=True, nbn_by_beds=True, drop_earlier=False):
    stays = [e for e in I.enrollments if e["type"] in (0, 8) or (e["type"] == 1 and not nbn_by_beds)]
    if drop_earlier:
        by = collections.defaultdict(list)
        for e in stays:
            by[(e["coc"], e["pid"])].append(e)
        gone = set()
        for lst in by.values():
            lst.sort(key=lambda e: (e["entry"], e["eid"]))
            for a, b in zip(lst, lst[1:]):
                if a["exit"] is None or a["exit"] > b["entry"]:
                    gone.add((a["coc"], a["eid"]))
        stays = [e for e in stays if (e["coc"], e["eid"]) not in gone]
    seen = collections.defaultdict(set)
    cnt = collections.Counter()
    for e in stays:
        end = e["exit"] if e["exit"] is not None else HI + dt.timedelta(days=1)
        a, b = max(e["entry"], LO), min(end, HI + dt.timedelta(days=1))
        for k in range((b - a).days):
            d = a + dt.timedelta(days=k)
            if dedup:
                seen[e["coc"]].add((e["pid"], d))
            else:
                cnt[(e["coc"], _q(d))] += 1
    if nbn_by_beds:
        for n in I.beds:
            if n["rt"] == 200 and LO <= n["date"] <= HI:
                if dedup:
                    seen[n["coc"]].add((n["pid"], n["date"]))
                else:
                    cnt[(n["coc"], _q(n["date"]))] += 1
    out = {}
    for c in COCS:
        if dedup:
            q = collections.Counter(_q(d) for p, d in seen[c])
            out[c] = [q.get(k, 0) for k in range(4)]
        else:
            out[c] = [cnt.get((c, k), 0) for k in range(4)]
    return out


def rrh_waits(I, chain="pid", rows="hoh"):
    """Average days from a household's entry into rapid re-housing to its move-in, over FY2026 move-ins, by Continuum
    and for the region (pooled). chain None, "hoh" or "pid" as in methods.rrh_waits; rows "all" reads a move-in
    from every enrollment carrying one."""
    rrh = [e for e in I.enrollments if e["type"] == 13]
    exits = collections.defaultdict(list)
    for e in rrh:
        if e["exit"] is not None:
            exits[(e["coc"], e["pid"], e["exit"])].append(e)
    kids = {(e["coc"], e["hh"]) for e in I.enrollments if e["rel"] == 2}
    waits = collections.defaultdict(list)
    for e in rrh:
        if e["move_in"] is None or not (LO <= e["move_in"] <= HI):
            continue
        if rows == "hoh" and e["rel"] != 1:
            continue
        fam = (e["coc"], e["hh"]) in kids
        start, cur = e["entry"], e
        while chain is not None:
            prev = [x for x in exits.get((cur["coc"], cur["pid"], start), ())
                    if x["eid"] != cur["eid"] and (chain == "pid" or (x["rel"] == 1 and cur["rel"] == 1))]
            if not prev:
                break
            cur = prev[0]
            start = cur["entry"]
        waits[(e["coc"], fam)].append((e["move_in"] - start).days)
    out, n = {}, {}
    for t in TYPES:
        for c in COCS:
            out[(c, t)] = sum(waits[(c, t)]) / len(waits[(c, t)])
            n[(c, t)] = len(waits[(c, t)])
        allw = [d for c in COCS for d in waits[(c, t)]]
        out[("ALL", t)] = sum(allw) / len(allw)
        n[("ALL", t)] = len(allw)
    out["n"] = n
    return out


def hround(v, nd=0):
    """Half away from zero, as a person rounds."""
    q = 10 ** nd
    return (int(v * q + 0.5) if v >= 0 else -int(-v * q + 0.5)) / q


class Golden:
    def __init__(self, target):
        I = self.I = Inputs(target)
        self.act = I.activity()
        self.roots = I.resolve()
        self.served = I.served(self.roots, self.act)
        self.counts = I.counts(self.served)
        self.amounts = allocation(self.counts)
        self.exact = allocation_exact(self.counts)
        self.nights = shelter_nights(I)
        self.rrh = rrh_waits(I)
