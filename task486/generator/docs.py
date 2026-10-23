"""task486 documents and workbooks: the Regional Fund Committee's allocation policy RF-2027, the State Data Linkage
Unit's method note with the run summary, the Partnership's guide to the integration extract, its project register, the
coalition's comparable-database export, the 2026 point-in-time count and housing inventory, and the data inventory."""
import collections
import csv
import datetime as dt

import docx
import openpyxl
from docx.shared import Pt, RGBColor
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import LETTER                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import inch                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import sim as S     # noqa: E402
import world as W    # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("LiberationSerif", "Regular"), ("LiberationSerif-Bold", "Bold"), ("LiberationSans", "Regular"),
              ("LiberationSans-Bold", "Bold"), ("LiberationSans-Italic", "Italic")):
    try:
        pdfmetrics.getFont(nm)
    except KeyError:
        pdfmetrics.registerFont(TTFont(nm, FONT_DIR + nm.split("-")[0] + "-" + f + ".ttf"))
INK = colors.HexColor("#1c1f22")
GREY = colors.HexColor("#5d636a")
NAVY = colors.HexColor("#1f3a5f")
SANS = ParagraphStyle("sans", fontName="LiberationSans", fontSize=9.4, leading=13.0, textColor=INK, spaceAfter=6)
CELL = ParagraphStyle("cell", fontName="LiberationSans", fontSize=8.0, leading=10.0, textColor=INK)
SMALL = ParagraphStyle("small", fontName="LiberationSans", fontSize=7.6, leading=9.8, textColor=GREY)


def _styles(accent):
    h1 = ParagraphStyle("h1", fontName="LiberationSans-Bold", fontSize=13.5, leading=17, textColor=accent,
                        spaceAfter=2)
    h1s = ParagraphStyle("h1s", fontName="LiberationSans", fontSize=9, leading=12, textColor=GREY, spaceAfter=10)
    h2 = ParagraphStyle("h2", fontName="LiberationSans-Bold", fontSize=10.2, leading=13.5, textColor=INK,
                        spaceBefore=7, spaceAfter=3, keepWithNext=1)
    return h1, h1s, h2


def _doc(path, title, author, footer):
    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(GREY)
        canvas.setFont("LiberationSans", 7.4)
        canvas.drawString(0.9 * inch, 0.55 * inch, footer)
        canvas.drawRightString(LETTER[0] - 0.9 * inch, 0.55 * inch, "Page %d" % doc.page)
        canvas.restoreState()
    d = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=0.9 * inch, rightMargin=0.9 * inch, topMargin=0.8 * inch,
                          bottomMargin=0.85 * inch, title=title, author=author, subject="", creator="", producer="")
    return d, frame


def _table(rows, widths, head="#e2e7ee"):
    t = Table([[Paragraph(str(c), CELL) for c in r] for r in rows], colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(head)),
                           ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#b9c0c4")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]))
    return t


# ------------------------------------------------------------------------------------------------ policy RF-2027
POLICY = [
    ("1. Purpose",
     ["The counties of Halvering, Melbray, Ondrava and Pyrenna pool the Halvering Valley Regional Homelessness Fund under "
      "their joint agreement of 2021. This policy sets how the Fund's 2027 amount is split among the region's four "
      "Continuums of Care and the Halvering Valley Coalition for Survivors."]),
    ("2. The amount and the split",
     ["The 2027 Fund is $18,750,000. The Halvering Valley Homeless Data Partnership keeps 4 per cent of it for "
      "administration and the regional data integration. The rest is split among the five recipients in proportion "
      "to the people each served in federal fiscal year 2026 (1 October 2025 to 30 September 2026). Each recipient's "
      "amount is rounded to the nearest dollar, and the Partnership's share takes up any dollar the rounding leaves "
      "over or short."]),
    ("3. People served",
     ["A person is counted once for the region if any record of theirs shows service in FY2026:",
      "(a) an enrollment in a Continuum's HMIS emergency shelter, safe haven, transitional housing, street outreach, "
      "rapid re-housing or permanent supportive housing project that was open on any day of the year, that is, "
      "entered on or before 30 September 2026 and not exited before 1 October 2025;",
      "(b) a stay in a coalition shelter on any night of the year;",
      "(c) a Wayfarer Street Health encounter in the year at which the patient's housing status was recorded as "
      "street, shelter or transitional housing; or",
      "(d) a contact by the City of Halvering Homeless Outreach Team in the year.",
      "Dates are the local dates the sources record."]),
    ("4. Recipient",
     ["Each person served is placed with one recipient: the Continuum of their latest HMIS enrollment open in the "
      "year, latest by entry date; for a person with no such enrollment, the coalition if they stayed in a coalition "
      "shelter in the year; otherwise the Continuum where their latest outreach contact or clinic encounter of the "
      "year took place. The outreach team works in Halvering, and the Partnership's register gives each clinic site's "
      "Continuum."]),
    ("5. One person across the sources",
     ["Each source holds people under its own identifiers. Records in different sources are matched by the State Data "
      "Linkage Unit, which compares the sources' hashed identifiers so that no source sees another's. The "
      "Partnership counts people from the unit's findings for its run of 8 October 2026 on the FY2025 and FY2026 "
      "records."]),
    ("6. Timetable",
     ["Halvering County, the Fund's fiscal agent, renews each recipient's contract for 2027 from 1 January 2027 at "
      "the amount this policy gives it. The Partnership prepares the table of people served and amounts for the "
      "Committee's meeting of 9 December 2026, so that the renewals are signed before the year begins."]),
]


def write_policy(path):
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    t = d.add_paragraph()
    r = t.add_run("Halvering Valley Regional Homelessness Fund")
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    t2 = d.add_paragraph()
    r2 = t2.add_run("Allocation policy RF-2027. Adopted by the Regional Fund Committee, 16 September 2026.")
    r2.italic = True
    r2.font.size = Pt(9.5)
    for head, paras in POLICY:
        h = d.add_paragraph()
        rh = h.add_run(head)
        rh.bold = True
        rh.font.size = Pt(11)
        for p in paras:
            q = d.add_paragraph(p)
            if p.startswith("("):
                q.paragraph_format.left_indent = Pt(18)
    d.core_properties.title = "Allocation policy RF-2027"
    d.save(path)


# ------------------------------------------------------------------------------------------------ the unit's note
def write_method_note(path, summary):
    h1, h1s, h2 = _styles(NAVY)
    d, frame = _doc(path, "Record linkage service: method note", W.LINKAGE_UNIT,
                    "%s, %s  |  Record linkage service, method note v3" % (W.LINKAGE_UNIT, W.LINKAGE_DEPT))
    story = [Paragraph("Record linkage service: method note", h1),
             Paragraph("%s, %s. Version 3, issued 3 June 2026. Run summary for the %s, run of 8 October 2026." %
                       (W.LINKAGE_UNIT, W.LINKAGE_DEPT, W.PARTNERSHIP), h1s)]
    paras = [
        ("The service",
         ["The unit links records that agencies hold about the same people without the agencies sharing identities. "
          "Each agency sends, for each of its records, its own record key and keyed hashes of the person's "
          "identifiers. The unit compares records from different sources and returns its findings on pairs of "
          "records, with the agreement of each identifier and the pair's score."]),
        ("Identifiers",
         ["First name; surname or surnames; year of birth; month and day of birth where the source holds them; sex; "
          "and the last four digits of the Social Security number where held. Names are normalised before hashing "
          "(case, accents, punctuation and spacing). A first name and a recognised nickname for it (Robert and Bob, "
          "Guadalupe and Lupe) agree partly, and a compound surname agrees partly with its first part. A source that "
          "holds only a year of birth sends no month or day, and those fields are then missing for every pair the "
          "record is in."]),
        ("Which records are compared",
         ["Two records are compared only when they come from different sources, share the year of birth and share "
          "the Soundex code of the first surname. Records without a year of birth are not compared."]),
        ("Scores and findings",
         ["Each identifier is coded A (agree), P (agree in part), D (disagree) or M (missing on either record), and "
          "the score is the sum of the weights below, year of birth included. A pair scoring 9.0 or more is returned "
          "as a LINK, one scoring from 4.0 to 8.9 as a NONLINK, and pairs under 4.0 are not returned.",
          "A LINK is the unit's finding that the two records very probably belong to the same person, and a NONLINK "
          "that they very probably do not. Each finding concerns one pair of records; the unit does not group "
          "records into people."]),
    ]
    for head, ps in paras:
        story.append(Paragraph(head, h2))
        for p in ps:
            story.append(Paragraph(p, SANS))
    wt = [["Identifier", "A", "P", "D", "M"]]
    for key, label in (("fn", "First name"), ("ln", "Surname"), ("by", "Year of birth"), ("bm", "Month of birth"),
                       ("bd", "Day of birth"), ("sex", "Sex"), ("ssn", "SSN, last four")):
        wts = S.World.WEIGHTS[key]
        wt.append([label] + ["%.1f" % wts[c] if c in wts else ("not used" if c == "M" and key in ("fn", "ln") else "")
                             for c in "APDM"])
    story.append(_table(wt, [1.6 * inch, 0.7 * inch, 0.7 * inch, 0.7 * inch, 0.8 * inch]))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Run summary: %s, run of 8 October 2026" % W.PARTNERSHIP, h2))
    story.append(Paragraph("Records received from the Partnership's FY2025 and FY2026 extract, by source, and the "
                           "pairs the run returned. Record keys are each source's own keys as sent.", SANS))
    rows = [["Source", "Records received", "Without a year of birth (not compared)"]]
    for s in W.SOURCES:
        rows.append([s, "{:,}".format(summary["received"][s]), "{:,}".format(summary["no_year"][s])])
    rows.append(["All sources", "{:,}".format(sum(summary["received"].values())),
                 "{:,}".format(sum(summary["no_year"].values()))])
    story.append(_table(rows, [1.6 * inch, 1.5 * inch, 2.6 * inch]))
    story.append(Spacer(1, 6))
    story.append(Paragraph("Pairs returned: {:,}, of which {:,} LINK and {:,} NONLINK.".format(
        summary["returned"], summary["links"], summary["returned"] - summary["links"]), SANS))
    story.append(Paragraph("Contact: %s, State Data Linkage Unit." % W.PEOPLE["linkage"], SMALL))
    d.build(story, onFirstPage=frame, onLaterPages=frame)


# ------------------------------------------------------------------------------------------------ the guide
GUIDE = """HALVERING VALLEY HOMELESS DATA PARTNERSHIP
Guide to the FY2026 integration extract
Kathleen Carlson, director of data. 14 October 2026.

The extract holds every record the region's sources hold for people served from 1 April 2024 to 30 September 2026, as
sent to the State Data Linkage Unit for its run of 8 October 2026, and the unit's findings from that run. Dates are local
dates (Pacific time) unless a field says otherwise.

FILES AND WHAT A ROW IS

hmis_client_fy2025_fy2026.csv
  One row per client record in a Continuum's HMIS. The four Continuums run separate HMIS implementations
  (SourceCoC: HAL Halvering, MEL Melbray, OND Ondrava, PYR Pyrenna). PersonalID identifies a person within a Continuum's
  HMIS: the HMIS lead's monthly duplicate review merges any two client records found to be the same person, so within
  one Continuum a person has one PersonalID. A person served in two Continuums has a record in each. Identifying fields
  are not in the extract; YearOfBirth and the data quality codes are.
hmis_enrollment_fy2025_fy2026.csv
  One row per project enrollment of a client (EnrollmentID), with the project, entry date, household, the
  client's relationship to the head of household and, in rapid re-housing and permanent supportive housing, the
  household's move-in date.
hmis_exit_fy2025_fy2026.csv
  One row per exit (one per enrollment that has ended), with the exit date and destination.
hmis_services_fy2025_fy2026.csv
  Bed nights at the night-by-night shelter (RecordType 200): one row per night a client slept there.
dvc_comparable_db_export.xlsx
  The Halvering Valley Coalition for Survivors' comparable database (domestic violence programs do not enter clients
  in HMIS). Sheet Clients: one client record per survivor, ClientKey identifying the survivor (intake looks up a
  returning survivor before opening a record). Birth year only, by the coalition's confidentiality policy. Sheet
  Stays: one row per shelter stay (StayStart, the first night; StayEnd, the day the stay ended).
wsh_hch_encounters_fy2025_fy2026.csv
  Wayfarer Street Health's health care for the homeless encounters: one row per encounter. PatientKey identifies a
  patient, the clinic registering each patient once. HousingStatus is the patient's housing status recorded at the
  encounter, in the HRSA categories (Street, Homeless Shelter, Transitional, Doubling Up, Unknown).
hot_field_app_contacts.jsonl
  The City of Halvering Homeless Outreach Team's field app: one line per contact. client_key identifies the client,
  the app keeping one key per client; contact_time is local time with its UTC offset; birth_year is the client's
  year of birth as the app holds it and dob_on_file whether it also holds the full date.
sdlu_linkage_pairs_run_20261008.csv
  The State Data Linkage Unit's findings: one row per pair of records the unit compared and returned. source_a and
  source_b are the sources (HMIS-HAL, HMIS-MEL, HMIS-OND, HMIS-PYR, DVC, WSH, HOT); record_a and record_b are the
  record keys as the source sent them. The unit's method note gives the agreement codes and weights.
partnership_project_register.xlsx
  The HMIS projects of the four Continuums, Wayfarer Street Health's sites and the outreach team's area.
pit_2026_counts_by_coc.csv, hic_2026_inventory.xlsx
  The region's 2026 point-in-time count (night of 28 January 2026) and housing inventory, as submitted to HUD.

RECORD KEYS
  HMIS PersonalIDs are each vendor's own: Halvering and Pyrenna numeric, Melbray with an M- prefix, Ondrava eight
  hexadecimal characters. The coalition's ClientKey (DV-), the clinic's PatientKey (numeric) and the outreach app's
  client_key (C-) are sent as held.

HMIS CONVENTIONS
  Project types follow the HMIS Data Standards: 0 emergency shelter (entry/exit), 1 emergency shelter (night by
  night), 2 transitional housing, 3 permanent supportive housing, 4 street outreach, 8 safe haven, 13 rapid
  re-housing.
  A client is in a project on each night from the entry date to the night before the exit date. An enrollment with
  no exit was still open when the extract was taken.
  At the night-by-night shelter (Halvering Winter Respite) each night a client sleeps there is a bed-night service.
  Under the Partnership's HMIS policy its enrollments are exited on the day after the last bed night once 30 days
  pass without one.
  HouseholdID groups the clients of one household's enrollment. RelationshipToHoH: 1 head of household, 2 child,
  3 spouse or partner, 4 other relative, 5 other.
  MoveInDate is the date the household moved into housing in a rapid re-housing or permanent supportive housing
  project; the HMIS Data Standards record it for the head of household.
  DOBDataQuality: 1 full date of birth reported, 2 approximate or partial, 8 client does not know, 9 client
  prefers not to answer, 99 not collected.

WHO TO ASK
  Extract and keys: Kathleen Carlson. The unit's findings: Joshua Bell, State Data Linkage Unit.
"""


def write_guide(path):
    with open(path, "w", encoding="utf-8", newline="\r\n") as fh:
        fh.write(GUIDE)


# ------------------------------------------------------------------------------------------------ register, PIT, HIC
TYPE_LABEL = {0: "Emergency Shelter - Entry Exit", 1: "Emergency Shelter - Night-by-Night", 2: "Transitional Housing",
              3: "PH - Permanent Supportive Housing", 4: "Street Outreach", 8: "Safe Haven",
              13: "PH - Rapid Re-Housing"}
OPENED = {1012: "2009-07-01", 1015: "2014-03-01", 1019: "2011-10-01", 1023: "2019-01-15", 1027: "2016-11-01",
          1031: "2010-01-01", 1034: "2021-05-01", 1038: "2012-06-01", 1042: "2008-09-01", 1046: "2017-02-01",
          1051: "2013-10-01", 1055: "2018-07-01", 1058: "2015-10-01", 1062: "2012-04-01", 1066: "2016-01-01",
          1069: "2022-08-01"}
DV_PROGRAMS = [("Hearthstone Emergency Shelter", "HAL", "ES"), ("Haven House", "MEL", "ES"),
               ("Bridges Transitional Program", "OND", "TH")]


def _sheet(ws, head, rows, widths):
    ws.append(head)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for k, wd in enumerate(widths):
        ws.column_dimensions[chr(65 + k)].width = wd


def write_register(path, world, site_name, site_coc):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HMIS projects"
    rows = []
    for p in sorted(world.projects, key=lambda p: (W.COCS.index(p["coc"]), p["id"])):
        note = "Open 1 November to 31 March each winter." if p["ptype"] == 1 else None
        rows.append([p["coc"], p["id"], p["name"], p["ptype"], TYPE_LABEL[p["ptype"]],
                     OPENED.get(p["id"], "2015-07-01"), None, note])
    _sheet(ws, ["SourceCoC", "ProjectID", "ProjectName", "ProjectType", "ProjectTypeName", "OperatingStartDate",
                "OperatingEndDate", "Notes"], rows, [10, 10, 40, 11, 34, 17, 16, 40])
    ws2 = wb.create_sheet("Clinic sites")
    _sheet(ws2, ["Site", "SiteName", "CoC"], [[k, site_name[k], site_coc[k]] for k in sorted(site_name)], [10, 62, 8])
    ws3 = wb.create_sheet("Outreach")
    _sheet(ws3, ["Team", "Area", "CoC"],
           [[W.OUTREACH, "City of Halvering, downtown and riverside zones Z01 to Z14", "HAL"]], [40, 52, 8])
    ws4 = wb.create_sheet("Coalition programs")
    _sheet(ws4, ["Program", "CoC", "Type"], [[n, c, t] for n, c, t in DV_PROGRAMS], [36, 8, 8])
    wb.save(path)
    return len(rows)


PIT_NIGHT = dt.date(2026, 1, 28)


def pit_rows(world):
    """The 2026 point-in-time count by Continuum (night of 28 January 2026), as each Continuum submitted it: sheltered
    people from its shelters' and the coalition's records for the night, unsheltered from the street count."""
    night = PIT_NIGHT
    sh = {c: collections.defaultdict(set) for c in W.COCS}
    for e in world.enrollments:
        if e["kind"] in ("ES", "SH", "TH") and e["entry"] <= night and (e["exit"] is None or e["exit"] > night):
            sh[e["coc"]][e["kind"]].add(e["pid"])
    for n in world.bed_nights:
        if n["date"] == night:
            sh[n["coc"]]["ES"].add(n["pid"])
    for s in world.dv_stays:
        if s["start"] <= night < s["end"]:
            name, coc, kind = DV_PROGRAMS[s["pid"] % 3]
            sh[coc][kind].add(("DV", s["pid"]))
    street = collections.Counter()
    for e in world.enrollments:
        if e["kind"] == "SO" and e["entry"] <= night and (e["exit"] is None or e["exit"] > night):
            if world.households[world.people[e["pid"]]["hh"]]["street"]:
                street[e["coc"]] += 1
    rows = []
    for c in W.COCS:
        es, th, sf = len(sh[c]["ES"]), len(sh[c]["TH"]), len(sh[c]["SH"])
        un = int(round(street[c] * (1.9 if c == "HAL" else 1.4)))
        rows.append([c, W.COC_NAME[c], es, th, sf, es + th + sf, un, es + th + sf + un])
    return rows


PIT_COLS = ["CoC", "CoCName", "Sheltered_ES", "Sheltered_TH", "Sheltered_SH", "Sheltered_Total", "Unsheltered",
            "Total"]


def write_hic(path, world):
    """The 2026 housing inventory count: beds and units by project on the night of 28 January 2026."""
    night = PIT_NIGHT
    occ = collections.defaultdict(set)
    peak = collections.Counter()
    for e in world.enrollments:
        if e["entry"] <= night and (e["exit"] is None or e["exit"] > night):
            occ[e["project"]].add(e["pid"])
    for n in world.bed_nights:
        if n["date"] == night:
            occ[1027].add(n["pid"])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventory"
    rows = []
    for p in sorted(world.projects, key=lambda p: (W.COCS.index(p["coc"]), p["id"])):
        if p["ptype"] == 4:
            continue
        n = len(occ[p["id"]])
        beds = int(5 * round((n * (1.08 if p["ptype"] in (0, 1, 8) else 1.12) + 4) / 5))
        units = int(round(beds / (1.0 if p["ptype"] in (0, 1, 8, 2) else 1.7)))
        seasonal = beds if p["ptype"] == 1 else 0
        rows.append([p["coc"], p["id"], p["name"], TYPE_LABEL[p["ptype"]], beds - seasonal, seasonal, units,
                     "Yes" if p["ptype"] != 1 else "Seasonal"])
    _sheet(ws, ["CoC", "ProjectID", "ProjectName", "ProjectType", "YearRoundBeds", "SeasonalBeds", "Units",
                "Availability"], rows, [8, 10, 40, 34, 14, 13, 8, 13])
    wb.save(path)
    return len(rows)


INVENTORY = []


def write_inventory(path, files, rows):
    """The Partnership's list of what the extract holds, as its data team keeps it."""
    meta = {
        "client": ("Four Continuums' HMIS (stacked)", "Client records, FY2025 and FY2026", "2026-10-07"),
        "enrollment": ("Four Continuums' HMIS (stacked)", "Project enrollments", "2026-10-07"),
        "exit": ("Four Continuums' HMIS (stacked)", "Exits", "2026-10-07"),
        "services": ("Four Continuums' HMIS (stacked)", "Bed nights at the night-by-night shelter", "2026-10-07"),
        "dvc": (W.COALITION, "Comparable database export: clients and shelter stays", "2026-10-06"),
        "vsm": (W.CLINIC, "Health care for the homeless encounters", "2026-10-06"),
        "hot": (W.OUTREACH, "Field app contacts", "2026-10-05"),
        "pairs": (W.LINKAGE_UNIT, "Pairwise findings, run of 8 October 2026", "2026-10-09"),
        "note": (W.LINKAGE_UNIT, "Method note v3 with the run summary", "2026-10-09"),
        "policy": (W.COMMITTEE, "Allocation policy RF-2027", "2026-09-16"),
        "guide": (W.PARTNERSHIP, "Guide to the FY2026 integration extract", "2026-10-14"),
        "register": (W.PARTNERSHIP, "Project register", "2026-10-02"),
        "pit": (W.PARTNERSHIP, "2026 point-in-time count by Continuum, as submitted", "2026-04-28"),
        "hic": (W.PARTNERSHIP, "2026 housing inventory count, as submitted", "2026-04-28"),
        "inventory": (W.PARTNERSHIP, "This list", "2026-10-14"),
    }
    INVENTORY[:] = []
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\r\n")
        w.writerow(["file", "from", "what", "rows", "received"])
        for k, name in files.items():
            frm, what, rec = meta[k]
            n = rows.get(k, "")
            w.writerow([name, frm, what, n, rec])
            INVENTORY.append((name, frm, what, n, rec))
    return len(INVENTORY)
