"""The golden deliverables, computed from the cut pack: the eligibility schedule workbook (the schedule, the vacancy
annex and notes) and the chart of every centre's change against the 8 per cent line. Returns figures.json's content:
every figure the write-up states, as it states it."""
import datetime as dt
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import openpyxl  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402

import golden as G  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
TARGET = os.path.join(os.path.dirname(HERE), "target")
DISTRICT_ORDER = ["Rookgate", "Dorrstead", "Glenskarn", "Ingmouth", "Neshdale", "Skelrigg"]


def r50(x):
    return int(round(x / 50.0)) * 50


def pct(x):
    return int(round(x * 100))


def compute():
    P, R = G.summary(TARGET)
    g22, g26 = R["golden"]
    ch = G.change(R["golden"])
    q = G.qualifying(R["golden"])
    V = G.vacancy(P)
    rows = []
    for c, (name, district) in P.centres.items():
        rows.append(dict(code=c, name=name, district=district, e22=r50(g22[c]), e26=r50(g26[c]), change=pct(ch[c]),
                         qualifies=c in q, raw=ch[c]))
    multi = {yr: [k for k, (T, ps) in P.retail(yr).items() if len(ps) > 1] for yr in ("22", "26")}
    comp = dict(
        multi_site_employers_2022=len(multi["22"]), multi_site_employers_2026=len(multi["26"]),
        multi_site_employees_2022=int(sum(P.retail("22")[k][0] for k in multi["22"])),
        multi_site_employees_2026=int(sum(P.retail("26")[k][0] for k in multi["26"])),
        retail_employers_2022=len(P.retail("22")), retail_employers_2026=len(P.retail("26")),
        retail_employees_2022=int(sum(T for T, ps in P.retail("22").values())),
        retail_employees_2026=int(sum(T for T, ps in P.retail("26").values())),
        centres_total_2022_thousands=int(round(sum(g22.values()) / 1000.0)),
        centres_total_2026_thousands=int(round(sum(g26.values()) / 1000.0)),
        qualifying=len(q))
    rung = {}
    for k in ("area", "rv", "count", "postcode"):
        s = G.qualifying(R[k])
        rung[k] = dict(n=len(s), added=[P.centres[c][0] for c in s if c not in q], dropped=[P.centres[c][0] for c in q if c not in s])
    return P, R, rows, comp, V, rung, q


def workbook(path, rows, comp, V):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Eligibility schedule"
    bold = Font(bold=True)
    thin = Side(style="thin", color="999999")
    ws.append(["Town Centre Renewal Fund 2027: eligibility schedule"])
    ws.append(["Workplace retail employment in the designated town centres, March 2022 and March 2026 (prospectus section 3)"])
    ws.append([])
    head = ["Code", "Town centre", "District", "March 2022", "March 2026", "Change (%)", "Meets the test"]
    ws.append(head)
    for r in rows:
        ws.append([r["code"], r["name"], r["district"], r["e22"], r["e26"], r["change"], "Yes" if r["qualifies"] else "No"])
    ws.append([])
    ws.append(["Employment is rounded to the nearest 50 and the change to the nearest whole per cent; the test is applied to "
               "unrounded figures. A centre meets the test when its workplace retail employment fell by 8 per cent or more."])
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"].font = Font(italic=True, color="555555")
    for c in ws[4]:
        c.font = bold
        c.border = Border(bottom=thin)
        c.alignment = Alignment(horizontal="center" if c.column > 3 else "left", wrap_text=True)
    fill = PatternFill("solid", fgColor="E3EEFB")
    for i, r in enumerate(rows, start=5):
        for col in (4, 5):
            ws.cell(i, col).number_format = "#,##0"
        ws.cell(i, 6).number_format = '+0;-0;0'
        ws.cell(i, 7).alignment = Alignment(horizontal="center")
        if r["qualifies"]:
            for col in range(1, 8):
                ws.cell(i, col).fill = fill
    for col, wd in zip("ABCDEFG", [7, 24, 12, 12, 12, 11, 15]):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A5"
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    v = wb.create_sheet("Vacancy annex")
    v.append(["Shop premises in each district's designated centres standing empty at 31 March"])
    v.append(["Valuation list entries described as shop and premises; empty when no one was in occupation on the day"])
    v.append([])
    v.append(["District", "Empty, 31 March 2022", "Of which empty 12 months or more", "Empty, 31 March 2026",
              "Of which empty 12 months or more"])
    t = [0, 0, 0, 0]
    for d in DISTRICT_ORDER:
        x = [V["22"][0][d], V["22"][1][d], V["26"][0][d], V["26"][1][d]]
        t = [a + b for a, b in zip(t, x)]
        v.append([d] + x)
    v.append(["Torvasay"] + t)
    v["A1"].font = Font(bold=True, size=13)
    v["A2"].font = Font(italic=True, color="555555")
    for c in v[4]:
        c.font = bold
        c.border = Border(bottom=thin)
        c.alignment = Alignment(wrap_text=True, horizontal="center" if c.column > 1 else "left")
    for c in v[v.max_row]:
        c.font = bold
        c.border = Border(top=thin)
    for col, wd in zip("ABCDE", [14, 14, 18, 14, 18]):
        v.column_dimensions[col].width = wd
    v.row_dimensions[4].height = 32
    n = wb.create_sheet("Notes")
    notes = [
        ("Prepared by", "Economic Statistics, Torvasay Statistics, for the Council of Ministers' meeting of 10 December 2026"),
        ("Employees", "PAYE schemes' employees paid in tax month 12 (6 March to 5 April), Revenue Office extracts of 16 May 2022 "
                      "and 18 May 2026; employers classified to SIC 2007 division 47, a company's schemes added together"),
        ("Premises", "The rating list in force on each 31 March (the 2017 list for 2022, the 2023 list for 2026) and the six "
                     "district councils' business rates accounts; a premises is in a centre when its list postcode is in the "
                     "Department's schedule, terminated codes included"),
        ("Where staff work", "A PAYE address is a correspondence address, so each employer's employees are placed at the "
                             "premises it occupied on 31 March. A single-site employer's employees are at its one premises; a "
                             "multi-site employer's are shared over its premises in proportion to floor area times the staff "
                             "per square metre of the premises' description, measured on the single-site retail employers in "
                             "that description and March"),
        ("Empty premises", "An account in an owner's name while a premises stands empty marks it empty: one on empty "
                           "property relief, or one whose liable party the companies register classes to letting its own "
                           "real estate or to developing building projects; a premises with no account on the day carried "
                           "no liability and was in occupation; an empty spell runs back through consecutive owners' "
                           "accounts, and counts as twelve months only where the premises was in the list throughout"),
        ("Coverage", "%d retail employers in March 2022 and %d in March 2026, of which %d and %d occupied more than one premises"
                     % (comp["retail_employers_2022"], comp["retail_employers_2026"], comp["multi_site_employers_2022"],
                        comp["multi_site_employers_2026"])),
        ("Rounding", "Employment to the nearest 50, so that no employer's payroll can be read from the schedule; changes to "
                     "the nearest whole per cent, from unrounded figures"),
    ]
    n.append(["Notes"])
    n["A1"].font = Font(bold=True, size=13)
    for a, b in notes:
        n.append([a, b])
    for i in range(2, n.max_row + 1):
        n.cell(i, 1).font = bold
        n.cell(i, 1).alignment = Alignment(vertical="top")
        n.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")
    n.column_dimensions["A"].width = 18
    n.column_dimensions["B"].width = 96
    wb.save(path)


def chart(path, rows, q):
    order = sorted(rows, key=lambda r: r["raw"])
    y = np.arange(len(order))
    vals = np.array([r["raw"] * 100 for r in order])
    qual = np.array([r["qualifies"] for r in order])
    fig, ax = plt.subplots(figsize=(8.6, 9.4), dpi=170)
    ink, muted, accent, neutral = "#0b0b0b", "#52514e", "#2a78d6", "#b7b5ad"
    ax.barh(y, vals, height=0.62, color=np.where(qual, accent, neutral), edgecolor="white", linewidth=1.2, zorder=3)
    ax.axvline(-8, color="#e34948", linewidth=1.6, linestyle=(0, (5, 3)), zorder=4)
    ax.text(-8.25, len(order) - 0.15, "8 per cent fall: the fund's test", color="#b3302f", fontsize=8.5, ha="right", va="bottom")
    ax.axvline(0, color=muted, linewidth=0.8, zorder=2)
    ax.set_yticks(y)
    ax.set_yticklabels([r["name"] for r in order], fontsize=8.2, color=ink)
    labels = []
    for yi, v, r in zip(y, vals, order):
        labels.append((v, ax.text(v + (0.35 if v >= 0 else -0.35), yi, "%+d%%" % r["change"], va="center",
                                  ha="left" if v >= 0 else "right", fontsize=7.4, color=muted, zorder=5)))
    ax.set_xlabel("Change in workplace retail employment, March 2022 to March 2026 (per cent)", fontsize=9, color=ink)
    lo, hi = np.floor(vals.min() / 5) * 5 - 4, np.ceil(vals.max() / 5) * 5 + 4
    ax.set_xlim(lo, hi)
    ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: "%+d" % v if v else "0"))
    ax.grid(axis="x", color="#e6e5e0", linewidth=0.7, zorder=0)
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.spines["bottom"].set_color(muted)
    ax.tick_params(axis="y", length=0)
    ax.tick_params(axis="x", colors=muted, labelsize=8)
    ax.set_ylim(-0.7, len(order) - 0.3 + 0.9)
    fig.text(0.015, 0.985, "%d of Torvasay's 31 town centres lost 8 per cent or more of their workplace retail jobs" % len(q),
             ha="left", va="top", fontsize=12.5, fontweight="bold", color=ink)
    fig.text(0.015, 0.957, "Employees of retail employers placed at the premises they work in, March 2022 to March 2026.\n"
             "Blue: the centres that meet the 2027 Town Centre Renewal Fund's test.", ha="left", va="top", fontsize=8.6,
             color=muted, linespacing=1.4)
    fig.text(0.015, 0.008, "Source: Torvasay Statistics from Revenue Office PAYE returns, Valuation Office rating lists and "
             "district councils' business rates accounts.", fontsize=7, color=muted)
    fig.tight_layout(rect=(0, 0.02, 1, 0.915))
    # a label that would sit on the 8 per cent line goes inside its bar's end instead
    fig.canvas.draw()
    inv = ax.transData.inverted()
    for v, tx in labels:
        (x0, _), (x1, _) = inv.transform(tx.get_window_extent())
        if x0 - 0.15 <= -8 <= x1 + 0.15:
            tx.set_x(v + 0.35)
            tx.set_ha("left")
            tx.set_color(ink)
    fig.savefig(path, metadata={"Software": None, "Title": "Workplace retail employment change by town centre",
                                "Author": "Torvasay Statistics"})
    plt.close(fig)


def write(gpath, w=None):
    P, R, rows, comp, V, rung, q = compute()
    workbook(gpath("workbook"), rows, comp, V)
    chart(gpath("chart"), rows, q)
    figs = dict(call=dict(qualifying=[P.centres[c][0] for c in q], n=len(q)),
                components=comp,
                schedule=[dict(code=r["code"], centre=r["name"], employment_2022=r["e22"], employment_2026=r["e26"],
                               change_pct=r["change"], qualifies="Yes" if r["qualifies"] else "No") for r in rows],
                asks={"ask_a": {d: {"2022": V["22"][0][d], "2026": V["26"][0][d]} for d in DISTRICT_ORDER},
                      "ask_b": {d: {"2022": V["22"][1][d], "2026": V["26"][1][d]} for d in DISTRICT_ORDER}},
                ladder=rung)
    # the conventions and totals the write-up states beside the graded figures
    figs["stated"] = dict(sic_division="47", employment_unit=50, test_clause="3.3",
                          centres_total_2022=comp["centres_total_2022_thousands"] * 1000,
                          centres_total_2026=comp["centres_total_2026_thousands"] * 1000,
                          empty_2022=sum(V["22"][0].values()), long_2022=sum(V["22"][1].values()),
                          empty_2026=sum(V["26"][0].values()), long_2026=sum(V["26"][1].values()),
                          smallest_qualifying_fall=round(-100 * max(r["raw"] for r in rows if r["qualifies"]), 1),
                          largest_other_fall=round(-100 * min(r["raw"] for r in rows if not r["qualifies"]), 1),
                          owner_class_letting="68209", owner_class_development="41100")
    return figs
