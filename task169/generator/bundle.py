"""Read the evidence pack back into the arrays the forecast runs on, the way an analyst conforms it: the legacy
export (local hours, two channels) and the platform export (UTC hours by direction) onto local clock hours of each
counted segment, capacities from the inventory's minimum widths and the guide's Table 4.3, the calendar from the
published term dates and the weather from the station record. Used by deliver.py and cracker.py; verify.py has
its own reader."""
import csv
import os
import re
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

import numpy as np
import openpyxl
import pdfplumber
import pyarrow.parquet as pq

LONDON = ZoneInfo("Europe/London")
YEARS = [2024, 2025, 2026]


def summer_days(y):
    d, out = date(y, 6, 1), []
    while d <= date(y, 8, 31):
        out.append(d)
        d += timedelta(days=1)
    return out


def read_register(target):
    wb = openpyxl.load_workbook(os.path.join(target, "counter_register.xlsx"), read_only=True)
    ws = wb["Register"]
    rows = list(ws.iter_rows(values_only=True))
    head = rows[0]
    return [dict(zip(head, r)) for r in rows[1:] if r[0]]


def read_inventory(target):
    wb = openpyxl.load_workbook(os.path.join(target, "cycleway_path_inventory_2026.xlsx"), read_only=True)
    rows = list(wb["Segments"].iter_rows(values_only=True))
    head = rows[0]
    return [dict(zip(head, r)) for r in rows[1:] if r[0]]


def read_bands(target):
    with pdfplumber.open(os.path.join(target, "cycle_design_guide_2023_section4.pdf")) as p:
        text = "\n".join(pg.extract_text() for pg in p.pages)
    bands = []
    for m in re.finditer(r"(\d\.\d) to under (\d\.\d) ([\d,]+)", text):
        bands.append((float(m.group(1)), float(m.group(2)), int(m.group(3).replace(",", ""))))
    m = re.search(r"(\d\.\d) and over ([\d,]+)", text)
    bands.append((float(m.group(1)), 99.0, int(m.group(2).replace(",", ""))))
    return bands


def capacity(bands, w):
    for lo, hi, cap in bands:
        if lo <= w < hi:
            return cap
    raise ValueError(w)


def read_calendar(target):
    hol, bank = [], set()
    with open(os.path.join(target, "school_term_dates_2023-2027.csv"), newline="") as fh:
        for r in csv.DictReader(fh):
            a = date.fromisoformat(r["first_day"])
            b = date.fromisoformat(r["last_day"])
            if r["type"] == "school holiday":
                hol.append((a, b))
            elif r["type"] == "bank holiday":
                bank.add(a)

    def dt(d):
        if d.weekday() >= 5 or d in bank:
            return "we"
        return "wk_hol" if any(a <= d <= b for a, b in hol) else "wk_term"
    return dt


def read_weather(target):
    wx = {}
    with open(os.path.join(target, "weather_station_daily_2024-01-01_2026-09-30.csv"), newline="") as fh:
        for r in csv.DictReader(fh):
            if r["tmax_c"] == "":
                continue
            wx[date.fromisoformat(r["date"])] = (float(r["tmax_c"]), float(r["rain_mm"]), float(r["sunshine_h"]),
                                                 float(r["tmin_c"]))
    return wx


def load(target, estimated_as_recorded=False, utc_as_local=False, nodata_as_zero=False):
    """flows[segment_ref][year] -> days x 24 local hours, NaN where not recorded; caps, cycleways, days, dts, weather.
    The keyword switches are the mishandlings checks.py runs against the golden (R2)."""
    reg = read_register(target)
    inv = read_inventory(target)
    bands = read_bands(target)
    seg_of_counter = {r["counter_id"]: r["segment_ref"] for r in reg}
    seg_of_legacy = {}
    for r in reg:
        seg_of_legacy[r["legacy_site_name"]] = r["segment_ref"]
    caps, cycleways = {}, {}
    for s in inv:
        if s["counter_id"]:
            caps[s["segment_ref"]] = capacity(bands, float(s["min_width_m"]))
            cycleways.setdefault(s["cycleway"], []).append(s["segment_ref"])
    days = {y: summer_days(y) for y in YEARS}
    pos = {y: {d: i for i, d in enumerate(days[y])} for y in YEARS}
    flows = {s: {y: np.full((len(days[y]), 24), np.nan) for y in YEARS} for s in caps}
    # legacy, summer 2024
    with open(os.path.join(target, "counts_summer2024_legacy.csv"), newline="", encoding="utf-8") as fh:
        lines = fh.read().splitlines()
    for r in csv.DictReader(lines[1:], delimiter=";"):
        seg = seg_of_legacy[r["Site"]]
        d = datetime.strptime(r["Date"], "%d/%m/%Y").date()
        h = int(r["Hour"][:2])
        if r["Ch1"] == "" or r["Ch2"] == "":
            continue
        flows[seg][2024][pos[2024][d], h] = float(r["Ch1"]) + float(r["Ch2"])
    # platform, summers 2025 and 2026
    t = pq.read_table(os.path.join(target, "counts_summer2025_2026.parquet")).to_pydict()
    acc = {}
    for cid, ts, dr, c, st in zip(t["counter_id"], t["timestamp_utc"], t["direction"], t["count"], t["status"]):
        loc = ts if utc_as_local else ts.astimezone(LONDON)
        d, h = loc.date(), loc.hour
        if d.year not in (2025, 2026) or d not in pos[d.year]:
            continue
        key = (seg_of_counter[cid], d.year, pos[d.year][d], h)
        ok = st == "ok" or (estimated_as_recorded and st == "estimated") or (nodata_as_zero and st == "no_data")
        a = acc.setdefault(key, [0.0, True])
        a[0] += c
        a[1] = a[1] and ok
    for (seg, y, i, h), (v, ok) in acc.items():
        if ok:
            flows[seg][y][i, h] = v
    if nodata_as_zero:
        for seg in flows:
            flows[seg][2024] = np.nan_to_num(flows[seg][2024], nan=0.0)
    dtf = read_calendar(target)
    dts = {y: [dtf(d) for d in days[y]] for y in YEARS}
    wx = read_weather(target)
    weather = {d: (v[0], v[1]) for d, v in wx.items()}
    return dict(flows=flows, caps=caps, cycleways=cycleways, days=days, dts=dts, weather=weather, wx=wx,
                bands=bands, register=reg, inventory=inv)
