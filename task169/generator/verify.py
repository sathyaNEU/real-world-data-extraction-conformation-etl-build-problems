"""task169 verify: an independent recomputation from the bundle alone (target/, figures.json, submission.md).

It shares no code with calc.py: its own readers for every format the pack ships, its own linking, exposure,
model and projection, and its own counts for the asks. Exit 1 on any mismatch."""
from __future__ import annotations

import csv
import datetime as dt
import json
import os
import re
import sys

import numpy as np
import openpyxl
import pdfplumber
import pyarrow.parquet as pq
import statsmodels.api as sm

TASK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BAD = []


def check(label, cond, detail=""):
    print("%-5s %-46s %s" % ("PASS" if cond else "FAIL", label, detail))
    if not cond:
        BAD.append(label)


def rows(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def d_mdy(s):
    return dt.datetime.strptime(s, "%m/%d/%Y").date() if s else None


def d_iso(s):
    return dt.date.fromisoformat(s[:10]) if s else None


def main(target):
    J = lambda f: os.path.join(target, f)
    figs = json.load(open(os.path.join(TASK, "generator", "figures.json")))
    sub = open(os.path.join(TASK, "submission.md"), encoding="utf-8").read()
    meth = " ".join((p.extract_text() or "") for p in pdfplumber.open(J("prison_population_projection_methodology_2026.pdf")).pages)
    meth = re.sub(r"\s+", " ", meth)
    check("PDF methodology read (control)", "730th day after it" in meth and "steps of five percentage points" in meth)
    table1 = {}
    for d, cs in re.findall(r"(Garrison|Ridgway|Linwood|Corliss) D\d ([A-Za-z, ]+?)(?= Garrison| Ridgway| Linwood| Corliss| 3\.)", meth):
        for c in cs.split(","):
            table1[c.strip().upper()] = d
    check("PDF methodology Table 1 read", len(table1) == 12, "%d counties" % len(table1))
    act = " ".join((p.extract_text() or "") for p in pdfplumber.open(J("act_26-114_enrolled.pdf")).pages)
    check("PDF act read (control)", "shall not recommit" in re.sub(r"\s+", " ", act))

    # ---- releases, admissions, arrests, cases
    rel = rows(J("doc_release_movements_2019-2026.csv"))
    norm = {"PAROLE": "PAR", "REPAROLE": "RPR", "MAX EXP": "EXP"}
    for r in rel:
        r["code"] = norm.get(r["movement_code"], r["movement_code"])
        r["rd"] = d_mdy(r["release_date"])
    adm = [r for r in rows(J("doc_admission_movements_2019-2026.csv")) if r["movement_type"] == "ADM"]
    for a in adm:
        a["t"] = d_iso(a["movement_datetime"])
    by_adm = {}
    for a in adm:
        by_adm.setdefault(a["doc_number"], []).append((a["t"], a["movement_reason"]))
    for v in by_adm.values():
        v.sort()
    by_rel = {}
    for r in rel:
        by_rel.setdefault(r["doc_number"], []).append(r["rd"])
    for v in by_rel.values():
        v.sort()
    xw = {r["sbi"].replace("-", ""): r["doc_number"] for r in rows(J("sbi_doc_number_crosswalk.csv"))}
    t = pq.read_table(J("sbi_arrests_released_persons_2019-2026.parquet")).to_pydict()
    grade_cls = lambda g: "F1" if g in ("FA", "FB") else "F2" if g.startswith("F") else "M"
    arrests = {}
    arr_rows = []
    for s, o, d, g in zip(t["sbi_no"], t["otn"], t["arrest_date"], t["top_charge_grade"]):
        doc = xw.get(s)
        if doc is None:
            continue
        arrests.setdefault(doc, []).append(d)
        arr_rows.append((doc, o, d, grade_cls(g)))
    for v in arrests.values():
        v.sort()
    death = {}
    for r in rows(J("vital_records_death_matches_2019-2026.csv")):
        if int(r["match_score"]) >= 80:
            death[r["doc_number"]] = min(death.get(r["doc_number"], dt.date(2099, 1, 1)), d_iso(r["date_of_death"]))
    cases = {}
    for r in sorted(rows(J("aoc_case_dispositions_2019-2026.csv")), key=lambda r: (r["docket_no"], r["entry_date"])):
        cases[r["otn"]] = r          # the latest entry per docket, one docket per tracking number
    wb = openpyxl.load_workbook(J("board_violation_actions_legacy_2019-2026.xlsx"), read_only=True)
    legacy = []
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        head = next(it)
        for row in it:
            legacy.append(dict(zip(head, row)))
    check("XLSX legacy board read", len(legacy) > 20000, "%d rows" % len(legacy))
    new = rows(J("board_case_decisions_2026.csv"))
    revoked = {r["arrest_otn"] for r in legacy if r["action"] == "REVOKE" and r["arrest_otn"]} | \
              {r["arrest_otn"] for r in new if r["decision"] == "REVOKED" and r["arrest_otn"]}

    cut = dt.date(2026, 9, 30)
    complete = cut - dt.timedelta(days=730)

    def custody(doc):
        out = []
        rl = by_rel.get(doc, [])
        for a, _ in by_adm.get(doc, []):
            nxt = next((x for x in rl if x > a), dt.date(2099, 1, 1))
            out.append((a, nxt))
        return out

    def first_adm(doc, rd):
        for a, why in by_adm.get(doc, []):
            if a > rd:
                return a, why
        return None, None

    def feat(age, priors, off, female, dist):
        return [1.0, age, np.log1p(priors), off == "PROPERTY", off == "DRUG", off == "PUBLIC ORDER", female,
                dist == "Ridgway", dist == "Linwood", dist == "Corliss"]

    bands = ((0, 90), (90, 180), (180, 365), (365, 730))
    Xr, Y, E, B = [], [], [], []
    Xl, yl = [], []
    for r in rel:
        if r["code"] not in ("PAR", "RPR") or r["rd"] > complete:
            continue
        doc, rd = r["doc_number"], r["rd"]
        age = (rd - d_mdy(r["birth_date"])).days / 365.25
        f = feat(age, int(r["prior_commitments"]), r["governing_offense"], r["sex"] == "F", table1[r["county_of_commitment"]])
        fa, why = first_adm(doc, rd)
        Xl.append(f)
        yl.append(1.0 if (fa is not None and (fa - rd).days <= 730 and why == "NEW COMMITMENT") else 0.0)
        cu = custody(doc)
        dd = death.get(doc)
        ar = arrests.get(doc, [])
        for k, (lo, hi) in enumerate(bands):
            a = rd + dt.timedelta(days=lo)
            b = rd + dt.timedelta(days=hi)
            if dd is not None and dd < b:
                b = dd
            if b <= a:
                continue
            ex = (b - a).days - sum(max(0, (min(b, e) - max(a, s)).days) for s, e in cu)
            if ex <= 0:
                continue
            n = sum(1 for x in ar if a <= x < b)
            Xr.append(f + [k == 1, k == 2, k == 3])
            Y.append(n)
            E.append(ex)
    Xr = np.array(Xr, float)
    m = sm.GLM(np.array(Y, float), Xr, family=sm.families.Poisson(), offset=np.log(np.array(E, float))).fit()
    mr1 = sm.Logit(np.array(yl), np.array(Xl, float)).fit(disp=0)
    # conversion by class, from arrests to 30 June 2024 the courts decided
    lim = dt.date(2024, 6, 30)
    stat = {c: [0, 0] for c in ("F1", "F2", "M")}
    lags = {c: [] for c in stat}
    off_of = {}
    for r in rel:
        off_of.setdefault(r["doc_number"], r["governing_offense"])
    mix = {}
    for doc, o, d, c in arr_rows:
        if d > lim:
            continue
        mix.setdefault(off_of.get(doc), {}).setdefault(c, 0)
        mix[off_of.get(doc)][c] += 1
        if o in revoked:
            continue
        cs = cases.get(o)
        isdoc = cs is not None and cs["sentence_type"] == "DOC"
        stat[c][0] += 1
        stat[c][1] += isdoc
        if isdoc:
            sd = d_iso(cs["sentence_date"])
            nc = [a for a, why in by_adm.get(doc, []) if why == "NEW COMMITMENT" and sd <= a <= sd + dt.timedelta(days=60)]
            if nc:
                lags[c].append((nc[0] - d).days)
    pc = {c: stat[c][1] / stat[c][0] for c in stat}
    # SF17: the evidence behind the courts' conversion, each count from its own reader
    ev_rev = [o for doc, o, d, c in arr_rows if d <= lim and o in revoked]
    np_rev = sum(1 for o in ev_rev if cases.get(o, {}).get("disposition") == "NP")
    oth = [(o, c) for doc, o, d, c in arr_rows if d <= lim and o not in revoked]
    np_oth = sum(1 for o, c in oth if cases.get(o, {}).get("disposition") == "NP")
    isdoc = lambda o: cases.get(o, {}).get("sentence_type") == "DOC"
    pr = {}
    for r in rel:
        if r["code"] in ("PAR", "RPR"):
            pr.setdefault(r["doc_number"], []).append((r["rd"], d_mdy(r["max_expiration_date"]) if r["max_expiration_date"] else None))
    for v in pr.values():
        v.sort(key=lambda z: z[0])
    off = {c: [0, 0] for c in stat}
    for doc, o, d, c in arr_rows:
        if d > lim:
            continue
        last = [z for z in pr.get(doc, []) if z[0] <= d]
        if last and last[-1][1] is not None and d >= last[-1][1]:
            off[c][0] += isdoc(o)
            off[c][1] += 1
    every = [(o, c) for doc, o, d, c in arr_rows if d <= lim]
    want = figs
    check("SF17 revoked arrests all nolle prossed", {"arrests": len(ev_rev), "nolle_prosequi": np_rev} == want["conv_revoked_arrests"]
          and np_rev == len(ev_rev), "%d of %d" % (np_rev, len(ev_rev)))
    check("SF17 the courts' DOC sentences by class", [stat["F1"][1], stat["F1"][0]] == want["conv_courts"]["felony_ab"]
          and [stat["F2"][1], stat["F2"][0]] == want["conv_courts"]["felony_ce"] and [stat["M"][1], stat["M"][0]] == want["conv_courts"]["misdemeanour"]
          and len(oth) == want["conv_courts"]["arrests"] and np_oth == want["conv_courts"]["nolle_prosequi"], json.dumps(stat))
    check("SF17 the arrests after the maximum date", off["F1"] == want["conv_after_max_date"]["felony_ab"]
          and off["F2"] == want["conv_after_max_date"]["felony_ce"] and off["M"] == want["conv_after_max_date"]["misdemeanour"]
          and sum(v[1] for v in off.values()) == want["conv_after_max_date"]["arrests"], json.dumps(off))
    check("SF17 every arrest", {"arrests": len(every), "doc": sum(isdoc(o) for o, c in every)} == want["conv_every_arrest"], "")
    lagv = {c: np.sort(lags[c]) for c in lags}
    proj = rows(J("doc_projected_releases_2027.csv"))
    ds = np.arange(730)
    band_of = np.digitize(ds, [90, 180, 365])
    tot = 0.0
    by = {}
    p1 = []
    for r in proj:
        rd = d_iso(r["projected_release_date"])
        age = (rd - d_mdy(r["birth_date"])).days / 365.25
        dist = table1[r["county_of_commitment"]]
        f = feat(age, int(r["prior_commitments"]), r["governing_offense"], r["sex"] == "F", dist)
        lam = np.array([np.exp(m.params @ np.array(f + [k == 1, k == 2, k == 3], float)) for k in range(4)])[band_of]
        mx = mix[r["governing_offense"]]
        nn = sum(mx.values())
        w = sum(mx.get(c, 0) / nn * pc[c] * np.searchsorted(lagv[c], 730 - ds, side="right") / len(lagv[c]) for c in pc)
        p = 1 - np.exp(-(lam * w).sum())
        tot += p
        by.setdefault(dist, []).append(p)
        p1.append(1 / (1 + np.exp(-(mr1.params @ np.array(f, float)))))
    rate = tot / len(proj)
    step = int(5 * round(100 * rate / 5))
    dist_rate = {d: np.mean(v) for d, v in by.items()}
    hi = max(dist_rate, key=dist_rate.get)
    lo = min(dist_rate, key=dist_rate.get)
    check("CALL rate step", step == figs["call"]["rate_step"], "%.3f per cent, step %d" % (100 * rate, step))
    check("CALL cohort size", len(proj) == figs["cohort_size"], "%d" % len(proj))
    check("DEFAULT the record's new-commitment model misses", int(5 * round(100 * np.mean(p1) / 5)) != step,
          "%.2f per cent" % (100 * np.mean(p1)))

    # ---- ask A: one recommitment per parolee per hearing date, rescinded by 30 June 2026 out
    close, fy0, fy1 = dt.date(2026, 1, 2), dt.date(2025, 7, 1), dt.date(2026, 6, 30)
    def as_date(v):
        return v.date() if isinstance(v, dt.datetime) else v
    resc = {r["case_no"] for r in legacy if r["action"] == "RESCIND" and as_date(r["action_date"]) <= fy1}
    orders = {}
    for r in legacy:
        d = as_date(r["action_date"])
        if r["action"] == "REVOKE" and d and fy0 <= d <= min(fy1, close) and r["case_no"] not in resc:
            k = (r["doc_number"], d)
            o = orders.setdefault(k, [r["office"], False])
            o[1] = o[1] or bool(r["arrest_otn"])
    for r in new:
        d = d_iso(r["decision_ts"])
        if r["decision"] == "REVOKED" and d and fy0 <= d <= fy1 and not (r["rescinded_date"] and d_iso(r["rescinded_date"]) <= fy1):
            k = (r["doc_number"], d)
            o = orders.setdefault(k, [r["office"], False])
            o[1] = o[1] or bool(r["arrest_otn"])
    A = {o: {"arrest": 0, "conditions": 0} for o in ("D1", "D2", "D3", "D4")}
    for off, arrest in orders.values():
        A[off]["arrest" if arrest else "conditions"] += 1
    check("ASK A every office and basis", all(A[o] == figs["a1_" + o] for o in A), json.dumps(A))
    # ---- ask B: parolee-days, each parolee's overlapping detainers merged
    moved = {(r["doc_number"], r["report_date"]) for r in new if r["opened_in"] == "LEGACY"}
    last = {}
    for r in legacy:
        if r["hold_lodged"]:
            k = (r["case_no"], as_date(r["hold_lodged"]))
            if k not in last or as_date(r["action_date"]) > as_date(last[k]["action_date"]):
                last[k] = r
    spans = []
    for r in last.values():
        s = as_date(r["hold_lodged"])
        if (r["doc_number"], as_date(r["report_date"]).isoformat()) in moved:
            continue
        e = as_date(r["hold_lifted"]) if r["hold_lifted"] else None
        if e is None:
            e = next((a for a, _ in by_adm.get(r["doc_number"], []) if a >= s), None)
        if e is None:
            continue
        spans.append((r["doc_number"], table1[r["hold_facility"].replace(" CJ", "").upper()], s, e))
    for r in new:
        if not r["hold_start"]:
            continue
        e = d_iso(r["hold_end"]) or cut + dt.timedelta(days=1)
        spans.append((r["doc_number"], table1[r["hold_facility"].replace(" County Jail", "").upper()], d_iso(r["hold_start"]), e))
    held = {}
    for doc, dist, s, e in spans:
        day = s
        while day < e:
            if dt.date(2025, 1, 1) <= day <= dt.date(2025, 12, 31):
                held.setdefault((doc, dist), set()).add(day)
            day += dt.timedelta(days=1)
    Bv = {d: {"H1": 0, "H2": 0} for d in ("Garrison", "Ridgway", "Linwood", "Corliss")}
    for (doc, dist), days in held.items():
        for day in days:
            Bv[dist]["H1" if day.month <= 6 else "H2"] += 1
    check("ASK B every district and half", all(Bv[d] == figs["b_" + d] for d in Bv), json.dumps(Bv))
    # ---- ask C: writ trips, ended by the return or by the release
    allm = rows(J("doc_admission_movements_2019-2026.csv"))
    ev = {}
    for m_ in allm:
        if m_["movement_type"] in ("WRO", "WRI", "ADM", "TRN"):
            ev.setdefault(m_["doc_number"], []).append((d_iso(m_["movement_datetime"]), {"ADM": 1, "TRN": 2, "WRI": 3, "WRO": 4}[m_["movement_type"]], m_["movement_type"], m_["location"]))
    for r in rel:
        ev.setdefault(r["doc_number"], []).append((r["rd"], 0, "REL", ""))
    Cv = {}
    nested = orphan = 0
    for doc, es in ev.items():
        op = None
        for t_, _, mt, loc in sorted(es):
            if mt == "WRO":
                nested += op is not None
                op = (t_, loc)
            elif mt == "WRI" and op is None:
                orphan += 1
            elif op is not None and mt in ("WRI", "REL"):
                day = op[0]
                while day < t_:
                    if day.year == 2025:
                        c = Cv.setdefault(op[1], {"H1": 0, "H2": 0})
                        c["H1" if day.month <= 6 else "H2"] += 1
                    day += dt.timedelta(days=1)
                op = None
    check("ASK C one writ trip open at a time per person, every WRI closing one", nested == 0 and orphan == 0,
          "%d WRO with a trip open, %d WRI with none" % (nested, orphan))
    check("ASK C every institution and half", all(Cv.get(i) == figs["c_" + i] for i in ("SCI-KTR", "SCI-FNW", "SCI-DNM", "SCI-LKV", "SCI-BRW")), json.dumps(Cv))
    # ---- the write-up states the figures
    for d in Bv:
        for q in Bv[d]:
            check("WRITE-UP carries %s %s" % (d, q), "{:,}".format(Bv[d][q]) in sub)
    check("WRITE-UP carries the call", "%d per cent" % step in sub)
    # ---- the goldens read back
    g = openpyxl.load_workbook(os.path.join(TASK, "golden", "returns_projection_2027.xlsx"))
    v = g["Projection 2027"]["D4"].value
    check("GOLDEN workbook carries the rate", abs(v - step / 100) < 1e-9)
    print("\nverify: %d failed" % len(BAD))
    return 1 if BAD else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(TASK, "target")))
