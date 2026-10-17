"""task154's independent verifier. It reads only the shipped files under target/ (and, to compare, the goldens and
generator/figures.json), imports no generator module, and recomputes on its own code path: the call and the three
lines TBS's bulletin feeds, the default path and the other rungs, the back-test against the Office's past estimates,
the revision record and the coal and LNG table (the lodgement lag read from the ports' loadings, not assumed).

    python3 task154/generator/verify.py [target_dir golden_dir]"""
from __future__ import annotations

import csv
import datetime as dt
import json
import math
import os
import re
import sys

import openpyxl
import pdfplumber

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)


def half_up(x, nd):
    q = 10 ** nd
    return math.floor(x * q + 0.5) / q


def month_seq(a, b):
    y, m = int(a[:4]), int(a[5:])
    out = []
    while (y, m) <= (int(b[:4]), int(b[5:])):
        out.append("%04d-%02d" % (y, m))
        y, m = (y, m + 1) if m < 12 else (y + 1, 1)
    return out


def prev_month(mo, k=1):
    y, m = int(mo[:4]), int(mo[5:])
    for _ in range(k):
        y, m = (y, m - 1) if m > 1 else (y - 1, 12)
    return "%04d-%02d" % (y, m)


def q_months(q):
    y, n = int(q[:4]), int(q[-1])
    return ["%04d-%02d" % (y, 3 * n - 2 + k) for k in range(3)]


def q_prev(q):
    y, n = int(q[:4]), int(q[-1])
    return "%dQ%d" % ((y, n - 1) if n > 1 else (y - 1, 4))


def q_seq(a, b):
    out, q = [], a
    while True:
        out.append(q)
        if q == b:
            return out
        y, n = int(q[:4]), int(q[-1])
        q = "%dQ%d" % ((y, n + 1) if n < 4 else (y + 1, 1))


def num(s):
    return float(s.replace(",", ""))


# ---------------------------------------------------------------------------------------------------- reading
def read_csv(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def read_prices(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for sheet in ("PPI inputs", "PPI outputs", "Export prices", "Import prices"):
        ws = wb[sheet]
        codes = [str(ws.cell(row=3, column=j).value) for j in range(2, ws.max_column + 1)]
        d = {c: {} for c in codes}
        for r in range(5, ws.max_row + 1):
            when = ws.cell(row=r, column=1).value
            if not isinstance(when, dt.datetime):
                continue
            for j, c in enumerate(codes, 2):
                d[c]["%04d-%02d" % (when.year, when.month)] = float(ws.cell(row=r, column=j).value)
        out[sheet] = d
    return out


def read_bulletin(path):
    with pdfplumber.open(path) as pdf:
        text = "\n".join(p.extract_text() for p in pdf.pages)
    tables = {}
    parts = re.split(r"\n(?=Table \d\. )", text)
    months = ["2026-03", "2026-04", "2026-05", "2026-06"]
    for part in parts:
        m = re.match(r"Table (\d)\. ", part)
        if not m:
            continue
        n = int(m.group(1))
        rows = {}
        for ln in part.split("\n")[2:]:
            if n in (1, 2):
                mm = re.match(r"(\d) .*? ([\d,]+\.\d) ([\d,]+\.\d) ([\d,]+\.\d) ([\d,]+\.\d)$", ln)
                if mm:
                    rows[mm.group(1)] = dict(zip(months, (num(x) for x in mm.groups()[1:])))
            elif n in (5, 6, 7):
                mm = re.match(r"([A-Z]{3}) .*? ([\d,]+\.\d+) ([\d,]+\.\d+) ([\d,]+\.\d+) ([\d,]+\.\d+)$", ln)
                if mm:
                    rows[mm.group(1)] = dict(zip(months, (num(x) for x in mm.groups()[1:])))
            elif n == 4:
                mm = re.match(r"([A-Za-z ]+?) ([\d,]+\.\d|-) ([\d,]+\.\d|-) ([\d,]+\.\d|-) ([\d,]+\.\d|-)$", ln)
                if mm and mm.group(1) != "Total":
                    v = [0.0 if x == "-" else num(x) for x in mm.groups()[1:]]
                    rows[mm.group(1)] = {"coal": v[0], "coal_kt": v[1], "lng": v[2], "lng_kt": v[3]}
        tables[n] = rows
    return tables, text


def read_db(path):
    db, order, basis = {}, [], {}
    for r in read_csv(path):
        kind = "cvm" if r["price_basis"].startswith("Chain") else "cp"
        db[(r["release_id"], kind, r["quarter"], r["series_id"])] = int(r["value_m"])
        if r["release_id"] not in basis:
            order.append((r["release_date"], r["release_id"]))
            basis[r["release_id"]] = r["release_date"]
    order.sort()
    return db, [rid for _, rid in order]


def read_annual(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Table 3"]
    names = {"Gross domestic product": "GDP", "Changes in inventories: non-farm business": "CIPI_NF",
             "Exports of goods": "X_G", "Imports of goods": "M_G"}
    cols = {names[ws.cell(row=4, column=j).value]: j for j in range(2, ws.max_column + 1)
            if ws.cell(row=4, column=j).value in names}
    out = {k: {} for k in cols}
    for r in range(5, ws.max_row + 1):
        q = ws.cell(row=r, column=1).value
        if q:
            for k, j in cols.items():
                out[k][q.replace(" ", "")] = ws.cell(row=r, column=j).value
    return out


def read_tracker(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Tracker"]
    base, track = {}, {}
    for r in range(6, ws.max_row + 1):
        code = ws.cell(row=r, column=2).value
        if code:
            base[code] = ws.cell(row=r, column=3).value
            v = ws.cell(row=r, column=4).value
            if isinstance(v, (int, float)):
                track[code] = v
    return base, track


# ---------------------------------------------------------------------------------------------------- the revaluation
def rebased(idx):
    y25 = [v for k, v in idx.items() if k.startswith("2025")]
    a = sum(y25) / len(y25)
    return {k: v / a for k, v in idx.items()}


def period_price(p, mo, T):
    """The index averaged over the T months before the end of month mo, each month's index holding over its month:
    the integral of the step function over [end - T, end] divided by T."""
    lo, tot, k, edge = -float(T), 0.0, mo, 0.0
    while edge > lo + 1e-12:
        share = min(1.0, edge - lo)
        tot += share * p[k]
        edge -= 1.0
        k = prev_month(k)
    return tot / float(T)


def volume(stock, p, mo, rule, T):
    if rule == "turnover":
        return stock / period_price(p, mo, T)
    if rule == "monthend":
        return stock / p[mo]
    if rule == "qavg":
        return stock / ((p[mo] + p[prev_month(mo)] + p[prev_month(mo, 2)]) / 3.0)
    if rule == "two":
        return stock / period_price(p, mo, 2.0)
    raise ValueError(rule)


def inventories(stocks, ratios, p25, q, rule):
    end, start = q_months(q)[-1], prev_month(q_months(q)[0])
    return sum(volume(stocks[c][end], p25[c], end, rule, ratios[c][end]) -
               volume(stocks[c][start], p25[c], start, rule, ratios[c][start]) for c in stocks)


# ---------------------------------------------------------------------------------------------------- main
def main(target=None, golden=None):
    target = target or os.path.join(TASK, "target")
    golden = golden or os.path.join(TASK, "golden")
    f = lambda name: os.path.join(target, name)
    names = sorted(os.listdir(target))
    pick = lambda pat: [n for n in names if re.match(pat, n)][0]
    prices = read_prices(f(pick(r"so_price_indexes_.*\.xlsx")))
    tabs, btext = read_bulletin(f(pick(r"tbs_merchandise_trade_business_stocks_.*\.pdf")))
    db, rel_order = read_db(f(pick(r"so_realtime_gdp_database\.csv")))
    annual = read_annual(f(pick(r"so_annual_national_accounts_.*\.xlsx")))
    base, track = read_tracker(f(pick(r"tarbert_gdp_tracker_.*\.xlsx")))
    portal = read_csv(f(pick(r"tbs_portal_monthly_.*\.csv")))
    bulk = read_csv(f(pick(r"tbs_trade_partner_sitc2_monthly_.*\.csv")))
    ports = read_csv(f(pick(r"harbours_board_bulk_loadings_.*\.csv")))
    countries = read_csv(f(pick(r"tbs_country_codes\.csv")))
    res, bad = {}, []

    def check(label, got, want, tol=1e-9):
        good = abs(got - want) <= tol
        print("%s  %s: %s against %s" % ("ok  " if good else "DIFF", label, got, want))
        if not good:
            bad.append(label)

    inds = list(tabs[5])
    p25 = {c: rebased(prices["PPI inputs"][c]) for c in inds}
    pout25 = {c: rebased(prices["PPI outputs"][c]) for c in inds}
    px25 = {s: rebased(prices["Export prices"][s]) for s in prices["Export prices"]}
    pm25 = {s: rebased(prices["Import prices"][s]) for s in prices["Import prices"]}
    q2 = ["2026-04", "2026-05", "2026-06"]
    x2 = sum(tabs[1][s][m] / px25[s][m] for s in tabs[1] for m in q2)
    m2 = sum(tabs[2][s][m] / pm25[s][m] for s in tabs[2] for m in q2)
    gdp1 = base.pop("GDP")

    def call(cipi_q2):
        con = {c: (-1.0 if c == "M_S" else 1.0) * (track[c] - base[c]) / gdp1 * 100.0 for c in track}
        con["X_G"] = (x2 - base["X_G"]) / gdp1 * 100.0
        con["M_G"] = -(m2 - base["M_G"]) / gdp1 * 100.0
        con["CIPI_NF"] = (cipi_q2 - base["CIPI_NF"]) / gdp1 * 100.0
        qq = sum(con.values())
        return qq, ((1.0 + qq / 100.0) ** 4 - 1.0) * 100.0, con

    stocks, ratios = tabs[5], tabs[7]
    cq2 = inventories(stocks, ratios, p25, "2026Q2", "turnover")
    q, ann, con = call(cq2)
    res["call"] = ann
    rungs = {r: call(inventories(stocks, ratios, p25, "2026Q2", r))[1] for r in ("monthend", "qavg", "two")}
    out_cipi = sum(stocks[c]["2026-06"] / pout25[c]["2026-06"] - stocks[c]["2026-03"] / pout25[c]["2026-03"]
                   for c in inds)
    print("verify: the call %.4f (quarterly %.4f); month-end %.3f, quarter average %.3f, two months %.3f, output "
          "indexes %.3f; business inventories %.1f, goods %.1f / %.1f" % (
              ann, q, rungs["monthend"], rungs["qavg"], rungs["two"], call(out_cipi)[1], cq2, x2, m2))
    # the March quarter's goods volumes reproduce the base from TBS's January to March over the rebased indexes
    pst = {}
    for r in portal:
        pst[(r["series_id"], r["month"])] = float(r["value"])
    q1m = ["2026-01", "2026-02", "2026-03"]
    xq1 = sum(pst[("TBS.EXP.S%s.SA" % s, m)] / px25[s][m] for s in tabs[1] for m in q1m[:2]) + \
        sum(tabs[1][s]["2026-03"] / px25[s]["2026-03"] for s in tabs[1])
    check("March-quarter goods exports rebuilt within $1m of the tracker's base", abs(xq1 - base["X_G"]), 0.0, tol=1.0)
    # back-test against the Office's latest estimates of past quarters
    last = rel_order[-1]
    pstk = {c: {} for c in inds}
    prat = {c: {} for c in inds}
    for (sid, m), v in pst.items():
        part = sid.split(".")
        if part[1] == "STK":
            pstk[part[2]][m] = v
        elif part[1] == "RAT":
            prat[part[2]][m] = v
    for c in inds:
        pstk[c]["2026-03"] = stocks[c]["2026-03"]
        prat[c]["2026-03"] = ratios[c]["2026-03"]
    errs = {"turnover": [], "monthend": []}
    for qq in q_seq("2019Q2", "2026Q1"):
        off = db[(last, "cvm", qq, "CIPI_NF")]
        for rule in errs:
            errs[rule].append(inventories(pstk, prat, p25, qq, rule) - off)
    rmse = {k: math.sqrt(sum(e * e for e in v) / len(v)) for k, v in errs.items()}
    print("verify: back-test over %d quarters, RMSE turnover %.2f, month-end %.2f ($m)" % (
        len(errs["turnover"]), rmse["turnover"], rmse["monthend"]))
    # Ask A: the advance estimate's revision record
    lat = annual["GDP"]
    growth = lambda lv, qq: ((lv[qq] / lv[q_prev(qq)]) ** 4 - 1.0) * 100.0
    revs = {}
    for qq in q_seq("2023Q2", "2026Q1"):
        adv = {k[2]: v for k, v in db.items() if k[0] == qq + "-A" and k[1] == "cvm" and k[3] == "GDP"}
        revs[qq] = growth(lat, qq) - growth(adv, qq)
    va = list(revs.values())
    ask_a = {qq: revs[qq] for qq in ("2025Q3", "2025Q4", "2026Q1")}
    ask_a["mean"] = sum(va) / len(va)
    ask_a["mar"] = sum(abs(v) for v in va) / len(va)
    # Ask B: the lodgement lag read from the ports, then coal and LNG to five markets
    tbs_t = {}
    for r in bulk:
        if r["flow"] == "Exports" and r["sitc_division"] in ("32", "34"):
            k = (r["month"], r["sitc_division"])
            tbs_t[k] = tbs_t.get(k, 0) + int(r["quantity"])
    port_t = {}
    for r in ports:
        if r["cargo"] in ("Coal", "LNG"):
            k = (r["month"], "32" if r["cargo"] == "Coal" else "34")
            port_t[k] = port_t.get(k, 0) + int(r["tonnes_loaded"])

    def best_lag(months_):
        gaps = {lag: sum(abs(tbs_t[(m, d)] / port_t[(prev_month(m, lag), d)] - 1.0) for m in months_
                         for d in ("32", "34")) for lag in (0, 1)}
        return min(gaps, key=gaps.get), gaps
    lag25, g25 = best_lag(month_seq("2025-02", "2025-11"))
    lag26, g26 = best_lag(month_seq("2026-01", "2026-05"))
    print("verify: TBS's coal and LNG tonnes follow the ports' loadings by %d month(s) in 2025 and %d in 2026" % (
        lag25, lag26))
    code_of = {}
    for r in countries:
        code_of[r["country"]] = r["code"]
    val = {}
    for r in bulk:
        if r["flow"] == "Exports" and r["sitc_division"] in ("32", "34"):
            val[(r["month"], r["country"], r["sitc_division"])] = int(r["value_000"]) / 1000.0
    june = {}
    for name, row in tabs[4].items():
        code = code_of.get(name) or [c for n, c in code_of.items() if n.startswith(name)][0]
        june[(code, "32")], june[(code, "34")] = row["coal"], row["lng"]
    ask_b = {}
    for d in ("32", "34"):
        tot = 0.0
        for c in ("JP", "CN", "KR", "IN", "TW"):
            v26 = val.get(("2026-04", c, d), 0.0) + val.get(("2026-05", c, d), 0.0) + june.get((c, d), 0.0)
            v25 = sum(val.get((m if lag25 == 0 else month_seq(m, "2026-12")[lag25], c, d), 0.0)
                      for m in ("2025-04", "2025-05", "2025-06"))
            ask_b["%s_%s" % (c, d)] = v26 - v25
            tot += v26 - v25
        ask_b["total_%s" % d] = tot
    # against figures.json and the goldens
    with open(os.path.join(TASK if target == os.path.join(TASK, "target") else os.path.dirname(target), "generator",
                           "figures.json"), encoding="utf-8") as fh:
        figs = json.load(fh)
    check("the call, annualised, one decimal", half_up(ann, 1), figs["call"]["annualised_pct"])
    check("the call, unrounded", ann, figs["call"]["unrounded"], tol=1e-9)
    check("quarterly growth, two decimals", half_up(q, 2), figs["call"]["quarterly_pct"])
    check("business inventories' contribution", half_up(con["CIPI_NF"], 2), figs["contributions_pp"]["CIPI_NF"])
    check("goods exports' contribution", half_up(con["X_G"], 2), figs["contributions_pp"]["X_G"])
    check("goods imports' contribution", half_up(con["M_G"], 2), figs["contributions_pp"]["M_G"])
    check("business inventories, June quarter ($m)", round(cq2, 1), figs["business_inventories_m"]["june_quarter_2026"])
    check("goods exports, June quarter ($m)", round(x2, 1), figs["goods_m"]["exports_q2"])
    check("goods imports, June quarter ($m)", round(m2, 1), figs["goods_m"]["imports_q2"])
    check("the default path rounds to 2.2", half_up(rungs["monthend"], 1), 2.2)
    check("the back-test RMSE, turnover ($m)", round(rmse["turnover"], 1), figs["backtest_rmse_m"]["turnover"])
    check("the back-test RMSE, month-end ($m)", round(rmse["monthend"], 1), figs["backtest_rmse_m"]["month-end"])
    for k, v in ask_a.items():
        check("Ask A %s" % k, half_up(v, 1 if k.startswith("20") else 2), figs["ask_a"][k])
    check("Ask B lag in 2025 (months)", lag25, 1)
    check("Ask B lag in 2026 (months)", lag26, 0)
    for k, v in ask_b.items():
        check("Ask B %s" % k, half_up(v, 1), figs["ask_b"][k])
    gw = openpyxl.load_workbook(os.path.join(golden, "gdp_preview_2026q2_tables.xlsx"), data_only=True)
    check("the tables workbook's call cell", gw["Call"]["B3"].value, half_up(ann, 1))
    with pdfplumber.open(os.path.join(golden, "gdp_preview_2026q2.pdf")) as pdf:
        ptxt = " ".join(p.extract_text() for p in pdf.pages)
    check("the preview's headline carries the call", float("June quarter GDP: we expect %.1f%% annualised" %
                                                           half_up(ann, 1) in ptxt), 1.0)
    if bad:
        raise SystemExit("verify: %d differences: %s" % (len(bad), "; ".join(bad)))
    print("verify: every figure recomputed from the shipped files on an independent path")
    return res


if __name__ == "__main__":
    if len(sys.argv) == 3:
        main(sys.argv[1], sys.argv[2])
    else:
        main()
