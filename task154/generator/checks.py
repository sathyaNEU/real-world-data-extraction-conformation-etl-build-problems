"""task154's assertions, run by build.py on every build: the input gates, the three rules on the generator's clean
truth (R1, R2 per mess item, R3a, R3b, R3c), every semantic fork row (SF1 to SF18), the ladder, corridor, grid and
partial applications, the back-test, both asks with every stop, hazard and device, the separation line, the realism
sweep (personas, round figures, repeated blocks, containers, dates, invented names), the load-bearing facts each in
exactly one file, the distractors' deletion test, the goldens against figures.json and the pair arithmetic.

Every check prints PASS or FAIL with its figures; one FAIL stops the build."""
from __future__ import annotations

import csv
import datetime as dt
import io
import itertools
import json
import math
import os
import re
import shutil
import tempfile
import zipfile

import numpy as np

import asks as A
import calc as C
import common as K
import docs as D
import national as NA
import partners as PT
import records as R
import tracker as TR
import world as WD

TASK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MONTHS = C.MONTHS
r1 = lambda x: K.round_half_up(x, 1)
r2 = lambda x: K.round_half_up(x, 2)
LO, HI = 2.75, 2.85          # the golden's bin


class Checks:
    def __init__(self):
        self.rows = []

    def __call__(self, label, cond, detail=""):
        self.rows.append((label, bool(cond), detail))
        return bool(cond)


def in_bin(x):
    return LO <= x < HI


# ---------------------------------------------------------------------------------------------------- pack text
def pdf_text(path):
    import pypdf
    rd = pypdf.PdfReader(path)
    return re.sub(r"\s+", " ", " ".join(p.extract_text() or "" for p in rd.pages))


def xlsx_text(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(ws.title)
        for row in ws.iter_rows(values_only=True):
            out.extend(str(v) for v in row if v is not None)
    return re.sub(r"\s+", " ", " | ".join(out))


def pack_text(folder):
    out = {}
    for f in sorted(os.listdir(folder)):
        p = os.path.join(folder, f)
        if f.endswith(".pdf"):
            out[f] = pdf_text(p)
        elif f.endswith(".xlsx"):
            out[f] = xlsx_text(p)
        else:
            with open(p, encoding="utf-8") as fh:
                out[f] = re.sub(r"\s+", " ", fh.read())
    return out


# ---------------------------------------------------------------------------------------------------- clean truth
def clean_bundle(W):
    """The generator's clean truth as the call needs it: the world's own published tables at full precision, read
    without any of the pack's mess (the bulletin's revised April and May, seasonally adjusted trade, every index
    rebased to its 2025 average), and the tracker as built."""
    B = R.provisional_bundle(W, W.BASE)
    B["track"] = dict(W.TRACK)
    B["base"] = {k: v for k, v in W.BASE.items() if k != "GDP"}
    B["gdp_q1"] = W.BASE["GDP"]
    return B


def edit_distance(a, b):
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


# the reference register for H21: real organisations and places a reader of this pack could confuse an invented name
# with (statistics offices, central and commercial banks, wire services and polls, ports and terminals that ship coal
# and LNG, industry bodies), normalised by H21's rule
REAL = [
    "Statistics New Zealand", "Stats NZ", "Australian Bureau of Statistics", "Statistics Canada", "Statistics Norway",
    "Statistics Sweden", "Statistics Netherlands", "Statistics Denmark", "Statistics Finland", "Statistics Korea",
    "Statistics Indonesia", "Statistics South Africa", "Office for National Statistics", "Bureau of Economic Analysis",
    "Census Bureau", "Central Statistics Office", "Statistics Bureau of Japan", "National Bureau of Statistics",
    "Reserve Bank of Australia", "Reserve Bank of New Zealand", "Bank of England", "Bank of Canada",
    "Westpac", "ANZ", "ASB", "BNZ", "Kiwibank", "Commonwealth Bank", "National Australia Bank", "Macquarie",
    "Barclays", "HSBC", "Lloyds", "NatWest", "Santander", "Standard Chartered", "Bendigo Bank", "Suncorp",
    "Tandem", "TSB", "Rabobank", "Heartland Bank", "Cooperative Bank",
    "Reuters", "Bloomberg", "Dow Jones", "Refinitiv", "FactSet", "Consensus Economics", "FocusEconomics",
    "Market News International", "Informa", "Harbour Asset Management", "Infometrics", "BERL", "NZIER",
    "Newcastle", "Gladstone", "Hay Point", "Dalrymple Bay", "Abbot Point", "Port Kembla", "Port Waratah",
    "Lyttelton", "Port Taranaki", "Port Chalmers", "Port Nelson", "Napier", "Tauranga", "Marsden Point",
    "Port Hedland", "Dampier", "Karratha", "Withnell Bay", "Barrow Island", "Curtis Island", "Darwin", "Onslow",
    "Port Botany", "Port Adelaide", "Port Augusta", "Port Arthur", "Port Alma", "Richards Bay", "Roberts Bank",
    "Ras Laffan", "Bontang", "Arun", "Bintulu", "Gorgon", "Wheatstone", "Ichthys", "Prelude", "Kitimat",
    "North Head", "Kerrisdale", "Barrow Point", "Point Barrow", "Kesterson", "Blackwater", "Kestrel",
    "Business NZ", "Australian Industry Group", "Confederation of British Industry", "Manufacturers Federation",
    "Customs Service", "Port of Tauranga", "Ports of Auckland", "CentrePort", "Ardrossan", "Mallacoota",
]
GENERIC = {"the", "of", "and", "bank", "economics", "board", "office", "statistics", "department", "industry",
           "federation", "port", "ports", "cape", "bay", "point", "head", "harbour", "harbours", "line", "trade",
           "business", "national", "service", "group", "new", "zealand", "australian", "australia", "central",
           "bureau", "reserve", "market", "news", "international", "asset", "management", "island", "north", "south",
           "east", "west", "saltpan", "poll"}


def distinctive(name):
    words = re.findall(r"[a-z]+", name.lower().replace("'", ""))
    return [w for w in words if w not in GENERIC]


# ---------------------------------------------------------------------------------------------------- run
def run(target, golden, Bt):
    W, B = Bt.W, Bt.bundle
    ok = Checks()
    files = sorted(os.listdir(target))
    root = os.path.dirname(os.path.abspath(target))
    with open(os.path.join(root, "metadata.json"), encoding="utf-8") as fh:
        meta = json.load(fh)
    text = pack_text(target)
    figs = Bt.figs
    g = C.golden(B)
    G = g["ann"]
    routes = C.routes(B)
    bt = C.backtest(B)
    fa, rec = A.ask_a(B)
    fb = A.ask_b(B)
    q2 = "2026Q2"

    # ------------------------------------------------------------------------------------------------ input gates
    fmts = sorted({f.rsplit(".", 1)[1] for f in files})
    rows = {x["path"]: x["rows"] for x in meta["files"]}
    ok("G1 ten or more files", len(files) >= 10, "%d files" % len(files))
    ok("G2 three or more formats", len(fmts) >= 3, ", ".join(fmts))
    big = [f for f, n in rows.items() if n and n >= 25000]
    ok("G3 a file of 25,000 or more rows", big, "; ".join("%s %d" % (f, rows[f]) for f in big))
    ok("G4 two or more distractors declared and shipped", len(meta["distractor_files"]) >= 2 and
       all(f in files for f in meta["distractor_files"]), ", ".join(meta["distractor_files"]))
    ok("G5 no file name names a distractor, decoy, issue, memo or draft",
       not any(re.search(r"distract|decoy|issue|memo|draft|note|answer|trap", f) for f in files))
    ok("G6 metadata.json lists every shipped file with source, date and licence",
       sorted(x["path"] for x in meta["files"]) == files and
       all(x.get("source") and x.get("date") and x.get("license") for x in meta["files"]))

    # ------------------------------------------------------------------------------------------------ R1
    heads = []
    for f in files:
        p = os.path.join(target, f)
        if f.endswith(".csv"):
            with open(p, encoding="utf-8") as fh:
                heads.append(fh.readline())
    xl = " ".join(text[f] for f in files if f.endswith(".xlsx") or f.endswith(".pdf"))
    banned = r"volume of stocks|valuation adjustment|replacement cost|holding gain|stockholding|turnover period|" \
             r"inventory volume|revalu"
    ok("R1a no shipped column or table holds the deciding quantity", not re.search(banned, " ".join(heads) + xl, re.I))
    qs_db = {k[2] for k in B["db"]}
    ok("R1b no June-quarter 2026 estimate ships: the database and annual accounts end at the March quarter",
       "2026Q2" not in qs_db and "2026Q2" not in B["latest"]["GDP"] and
       all(c not in B["track"] for c in ("X_G", "M_G", "CIPI_NF")), "database quarters to %s" % max(qs_db))
    moved = C.call(B, C.cipi(B, q2, Tfun=lambda c, m: B["ratio"][c][MONTHS[m]] * 1.1))["ann"] - G
    B_p = dict(B)
    B_p["p25"] = {c: v * np.where(np.arange(len(v)) >= C.m_idx("2026-04"), 1.01, 1.0) for c, v in B["p25"].items()}
    moved_p = C.golden(B_p)["ann"] - G
    ok("R1c the golden's decisive computation reads TBS's book values, ratios and input price indexes",
       abs(moved) > 0.05 and abs(moved_p) > 0.05, "ratios x1.1 move it %+.3f; inputs +1%% from April %+.3f" %
       (moved, moved_p))
    ok("R1d the back-test's raw material ships: monthly stocks, cost of sales and ratios from January 2019 for "
       "eight industries, the Office's non-farm inventories for 28 quarters",
       all(len([m for m in B["portal_stocks"][c] if m >= "2019-01"]) == 89 for c in K.IND_CODES) and
       len(B["office_cipi"]) >= 28, "%d quarters" % len(B["office_cipi"]))

    # ------------------------------------------------------------------------------------------------ R2
    Bp = dict(B)
    for key in ("stocks", "sales", "ratio", "xsa", "msa"):
        Bp[key] = {k: dict(v) for k, v in B[key].items()}
    for c in K.IND_CODES:
        for m in ("2026-04", "2026-05"):
            Bp["stocks"][c][m] = B["portal_stocks"][c][m]
            Bp["ratio"][c][m] = B["portal_ratio"][c][m]
    for s in K.SITC_CODES:
        for m in ("2026-04", "2026-05"):
            Bp["xsa"][s][m] = B["portal_xsa"][s][m]
            Bp["msa"][s][m] = B["portal_msa"][s][m]
    portal_call = C.golden(Bp)["ann"]
    ok("R2.1 SF14 April and May from the 24 June portal pull instead of the bulletin's revisions: same call",
       in_bin(portal_call) and portal_call != G, "%.3f" % portal_call)
    stale = "2026Q1-T"
    stale_gdp = B["db"][(stale, "Chain", "2026Q1", "GDP")]
    stale_call = C.call(B, C.cipi(B, q2), gdp_q1=stale_gdp)["ann"]
    ok("R2.2 SF16 the 26 June download's March-quarter GDP as the bulletin lines' denominator: same call",
       in_bin(stale_call), "%.3f (GDP %d against %d)" % (stale_call, stale_gdp, B["gdp_q1"]))
    same_lines = all(B["db"][(stale, "Chain", q, c)] == B["latest"][c][q]
                     for q in NA.QUARTERS[NA.QI["2019Q1"]:] for c in ("X_G", "M_G", "CIPI_NF"))
    ok("R2.3 SF7 the 21 July update left goods trade and non-farm inventories where the 26 June download has them",
       same_lines and all(B["base"][c] == B["latest"][c]["2026Q1"] for c in ("X_G", "M_G", "CIPI_NF")))
    gq1 = {f: C.goods(B, "2026Q1", f) for f in ("x", "m")}
    B19 = dict(B)
    B19["px25"] = {s: np.array([B["px"][s].get(m, np.nan) for m in MONTHS]) / 100.0 for s in K.SITC_CODES}
    B19["pm25"] = {s: np.array([B["pm"][s].get(m, np.nan) for m in MONTHS]) / 100.0 for s in K.SITC_CODES}
    g19 = {f: C.goods(B19, "2026Q1", f) for f in ("x", "m")}
    Bdec = dict(B)
    dec = C.m_idx("2025-12")
    Bdec["px25"] = {s: v / v[dec] for s, v in B["px25"].items()}
    Bdec["pm25"] = {s: v / v[dec] for s, v in B["pm25"].items()}
    gdec = {f: C.goods(Bdec, "2026Q1", f) for f in ("x", "m")}
    err = lambda d: max(abs(d["x"] / B["base"]["X_G"] - 1), abs(d["m"] / B["base"]["M_G"] - 1))
    ok("R2.4 SF6 indexes rebased to their 2025 average reproduce the March quarter's goods volumes; at 2019 = 100 "
       "or December 2025 = 1 they miss them (loud)",
       abs(gq1["x"] - B["base"]["X_G"]) < 1.0 and abs(gq1["m"] - B["base"]["M_G"]) < 1.0 and err(g19) > 0.05 and
       err(gdec) > 0.005, "rebased %.1f/%.1f against %d/%d; 2019 base misses by %.1f%%, December by %.1f%%" % (
           gq1["x"], gq1["m"], B["base"]["X_G"], B["base"]["M_G"], 100 * err(g19), 100 * err(gdec)))
    Bor = dict(B)
    Bor["xor"], Bor["mor"] = B["portal_xor"], B["portal_mor"]
    gor = {"x": C.goods(Bor, "2026Q1", "x", src="xor"), "m": C.goods(Bor, "2026Q1", "m", src="mor")}
    ok("R2.5 SF10 original-terms trade misses the March quarter's goods volumes (loud)", err(gor) > 0.01,
       "misses by %.1f%%" % (100 * err(gor)))
    x4 = C.golden(B)["q"] * 4.0
    ok("R2.6 SF3 the quarterly change times four in place of compounding: same call", in_bin(x4), "%.3f" % x4)
    rel_ids = list(B["releases"])
    before = rel_ids[:rel_ids.index(R.LAYOUT_CHANGE)]
    after = rel_ids[rel_ids.index(R.LAYOUT_CHANGE):]
    has = lambda rid, s: any(k[0] == rid and k[3] == s for k in B["db"])
    ok("R2.7 the database's 2024 layout change leaves GDP under one series id in every release",
       all(has(r_, "GDP") for r_ in rel_ids) and has(before[-1], "GFCF_BUS") and not has(after[0], "GFCF_BUS") and
       has(after[0], "GFCF_ME"), "%d releases, layout change at %s" % (len(rel_ids), R.LAYOUT_CHANGE))
    add_ok = True
    for rid in (before[-1], after[0], rel_ids[-1]):
        series = R.OLD_SERIES if rid in before else R.NEW_SERIES
        for q in ("2024Q1", "2025Q1"):
            if (rid, "Chain", q, "GDP") not in B["db"]:
                continue
            parts = [s for s in series if s != "GDP"]
            tot = sum(NA.SIGN[s] * B["db"][(rid, "Chain", q, s)] for s in parts)
            add_ok &= abs(tot - B["db"][(rid, "Chain", q, "GDP")]) <= len(parts)
    ok("R2.8 in every layout the components add to GDP within whole-$m rounding", add_ok)
    Bnb = {k: v for k, v in B.items() if k not in ("bulk", "ports", "countries")}
    ok("R2.9 the bulk download (in $000), the ports' loadings and the country codes are off the call's path",
       C.golden(Bnb)["ann"] == G)
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(target, R.F_TRACKER), data_only=True)
    ws = wb["Tracker"]
    codes = [ws.cell(row=r_, column=2).value for r_ in range(6, ws.max_row + 1) if ws.cell(row=r_, column=2).value]
    ok("R2.10 the tracker's lines map one to one onto the Office's components",
       len(codes) == len(set(codes)) == 14 and set(codes) == set(B["latest"]), ", ".join(codes))
    rebuilt = routes["the March quarter rebuilt the same way"][1]
    ok("R2.11 SF8 the March quarter's inventories rebuilt by the same revaluation instead of its published change: "
       "same call", in_bin(rebuilt), "%.3f" % rebuilt)

    # ------------------------------------------------------------------------------------------------ R3
    CB = clean_bundle(W)
    c_def = C.call(CB, C.cipi(CB, q2, "monthend"))["ann"]
    c_gold = C.call(CB, C.cipi(CB, q2))["ann"]
    ok("R3a the analyst test on the clean truth: month-end deflation, every filed rule applied, files 2.2",
       r1(c_def) == 2.2 and not in_bin(c_def), "%.3f" % c_def)
    true_q2 = WD.true_cipi(W, q2)
    turn_q2 = C.cipi(CB, q2)
    m1_, m0_ = WD.MI["2026-06"], WD.MI["2026-03"]
    exact = sum(W.B[c][m1_] / K.window_avg(W.P[c], m1_, W.T[c][m1_]) - W.B[c][m0_] / K.window_avg(W.P[c], m0_, W.T[c][m0_])
                for c in K.IND_CODES)
    ok("R3b the decisive step alone on the clean truth files the golden, and recovers the physical change in stocks",
       r1(c_gold) == 2.8 and in_bin(c_gold) and abs(exact - true_q2) < 1e-6 and abs(turn_q2 - true_q2) < 10.0 and
       abs(C.cipi(CB, q2, "monthend") - true_q2) > 100.0,
       "%.3f; physical %.1f; the revaluation %.1f at the indexes' full precision, %.1f as published; month-end %.1f ($m)"
       % (c_gold, true_q2, exact, turn_q2, C.cipi(CB, q2, "monthend")))
    lower = {"quarter average": C.cipi(CB, q2, "qavg"), "two months": C.cipi(CB, q2, Tfun=lambda c, m: 2.0),
             "June end only": C.cipi(CB, q2, end="turnover", start="monthend"),
             "output price indexes": C.cipi(CB, q2, "output")}
    lw = {k: C.call(CB, v)["ann"] for k, v in lower.items()}
    ok("R3b each lower rung's single correction on the clean truth stays off the golden",
       all(not in_bin(v) and abs(v - c_gold) >= 0.1 for v in lw.values()),
       "; ".join("%s %.3f" % kv for kv in lw.items()))
    # R3c: no plant
    series = {}
    with open(os.path.join(target, R.F_PORTAL), encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            series.setdefault(row["series_id"], []).append(float(row["value"]))
    runs = [sid for sid, v in series.items() if any(v[i] == v[i + 1] == v[i + 2] == v[i + 3] for i in range(len(v) - 3))
            and not sid.startswith("TBS.RAT")]
    blocks = []
    for sid, v in series.items():
        if sid.startswith("TBS.RAT"):
            continue            # two-decimal ratios repeat by nature
        arr = np.array(v)
        for lag in range(1, 37):
            for i in range(len(arr) - lag - 5):
                if np.array_equal(arr[i:i + 6], arr[i + lag:i + lag + 6]):
                    blocks.append(sid)
                    break
            else:
                continue
            break
    vals = {sid: tuple(v) for sid, v in series.items()}
    # mineral fuel exports carry no seasonal pattern, so TBS publishes the same values for both adjustments
    same_adj = [sid for sid in vals if sid.endswith(".OR") and vals[sid] == vals.get(sid[:-3] + ".SA")]
    dup = len(vals) - len(set(vals.values())) - len(same_adj)
    rev = [a for a in vals for b in vals if a != b and vals[a] == tuple(reversed(vals[b]))]
    ytd = []
    months_p = R.PORTAL_MONTHS
    for sid, v in series.items():
        jan_drop = all(v[i] < v[i - 1] for i, m in enumerate(months_p) if m.endswith("-01") and i)
        mono = all(v[i] >= v[i - 1] for i, m in enumerate(months_p) if i and not m.endswith("-01"))
        if jan_drop and mono:
            ytd.append(sid)
    ok("R3c no copied, backfilled, swapped, reversed or year-to-date block in TBS's monthly series",
       not runs and not blocks and not dup and not rev and not ytd and same_adj == ["TBS.EXP.S3.OR"],
       "%d series; runs %s; blocks %s; duplicates %d; reversed %s; ytd %s; unadjusted as adjusted %s" % (
           len(series), runs[:3], blocks[:3], dup, rev[:2], ytd[:2], same_adj))
    docs_ = [f for f in files if f.endswith(".pdf") or f.endswith(".txt")]
    tracker_total = any(isinstance(ws.cell(row=r_, column=4).value, (int, float))
                        for r_ in range(6, ws.max_row + 1) if ws.cell(row=r_, column=2).value == "GDP")
    poll = json.load(open(os.path.join(target, R.F_POLL), encoding="utf-8"))
    pm = poll["summary"]["median"]
    ok("R3c no memo, draft or briefing: the documents are the Office's guide and notice and TBS's bulletin; the tracker "
       "states no total; the poll (a declared distractor, closed before the bulletin) is off both the golden and the "
       "default path",
       sorted(docs_) == sorted([R.F_GUIDE, R.F_NOTICE, R.F_BULLETIN]) and not tracker_total and
       R.F_POLL in meta["distractor_files"] and not (LO <= pm < HI) and pm != r1(routes["month-end"][1]) and
       poll["closed"][:10] < "2026-07-28", "poll median %.1f, closed %s" % (pm, poll["closed"]))

    # ------------------------------------------------------------------------------------------------ SF rows
    notice = text[R.F_NOTICE]
    guide = text[R.F_GUIDE]
    bull = text[R.F_BULLETIN]
    ok("SF1 the advance estimate is the Office's release of 30 July 2026",
       "advance estimate of gross domestic product for the June quarter 2026 will be published as scheduled, at "
       "10:45 on Thursday 30 July 2026" in notice and "30 July 2026" in guide)
    ok("SF2 the headline is the change from the previous quarter", "percentage change in the volume of GDP from the "
       "previous quarter" in guide)
    ok("SF3 compounding filed; times four lands the same call", "((1 + q/100) to the fourth power, less 1) times 100" in
       guide and r1(x4) == r1(G))
    ok("SF4 the June quarter is April to June 2026 in the bulletin", all(m in B["bulletin"]["stocks"]["WME"] for m in
                                                                          ("2026-04", "2026-05", "2026-06")))
    Bmay = dict(B)
    for key in ("stocks", "ratio", "xsa", "msa"):
        Bmay[key] = {k: dict(v) for k, v in B[key].items()}
        for k in Bmay[key]:
            Bmay[key][k]["2026-06"] = B[key][k]["2026-05"]
    may_t = C.golden(Bmay)["ann"]
    may_m = C.call(Bmay, C.cipi(Bmay, q2, "monthend"))["ann"]
    ok("SF5 the notice puts the June bulletin in the estimate; May carried into June misses",
       "The advance estimate will incorporate the June results." in notice and not in_bin(may_t) and
       not in_bin(may_m), "May carried: %.3f turnover, %.3f month-end" % (may_t, may_m))
    ok("SF6 2025 prices: only rebased indexes reproduce the March quarter (R2.4)", "The quarters of 2026 are at 2025 "
       "prices." in guide and abs(gq1["x"] - B["base"]["X_G"]) < 1.0)
    ok("SF7 the March-quarter base from the 26 June download lands the same call (R2.2, R2.3)",
       in_bin(stale_call) and same_lines)
    ok("SF8 the advance estimate revises no earlier quarter", "The advance estimate revises no earlier quarter, though "
       "each March quarter's advance estimate re-expresses the whole series" in guide)
    ok("SF9 goods are deflated by SITC section; the price workbook carries no total goods index",
       "export and import price indexes, by SITC section" in guide and
       set(B["px"]) == set(K.SITC_CODES) and set(B["pm"]) == set(K.SITC_CODES))
    ok("SF10 seasonally adjusted trade reproduces the March quarter, original terms do not (R2.5)",
       "Estimates are seasonally adjusted." in guide and err(gor) > 0.01)
    cover = bt["turnover"]["mean"]
    cov_call = C.call(B, C.cipi(B, q2) - cover)["ann"]
    ok("SF11 the Office's estimate for businesses outside the survey averages near zero; carrying its mean lands the "
       "same call", abs(cover) < 3.0 and abs(bt["turnover"]["last"]) < 5.0 and in_bin(cov_call),
       "mean %.1f, March quarter %.1f, call %.3f" % (cover, bt["turnover"]["last"], cov_call))
    outp = routes["output price indexes"][1]
    ok("SF12 the sources table names the producer price indexes of inputs; output indexes miss",
       "producer price indexes, inputs to each industry" in guide and not in_bin(outp), "%.3f" % outp)
    ratio_ok = all(abs(B["stocks"][c][m] / B["sales"][c][m] - B["ratio"][c][m]) < 0.006
                   for c in K.IND_CODES for m in R.BULLETIN_MONTHS)
    ok("SF13 the ratio is stocks over the month's cost of sales, and the bulletin's columns obey it; no selling value "
       "ships", "The stock-to-sales ratio is stocks at the end of the month divided by the month's cost of sales" in
       bull and ratio_ok and not any(".SAL." in s or "selling" in s.lower() for s in series))
    ok("SF14 the bulletin's April and May against the portal's: same call (R2.1)", in_bin(portal_call))
    ok("SF15 end-March stocks are the same in the bulletin and the portal",
       all(B["bulletin"]["stocks"][c]["2026-03"] == B["portal_stocks"][c]["2026-03"] for c in K.IND_CODES) and
       "March and earlier months are unchanged." in bull)
    ok("SF16 the 26 June download's March-quarter GDP as the denominator: same call (R2.2)", in_bin(stale_call))
    ok("SF17 rounding: the unrounded rate is a quarter of a bin or more from its rounding edge",
       K.bin_clearance(G, 0.1) >= 0.25 and "Releases give the headline to one decimal place, from the unrounded "
       "estimates." in guide, "%.4f, %.2f of a bin" % (G, K.bin_clearance(G, 0.1)))
    other_q2 = [f for f in (R.F_DB, R.F_ANNUAL, R.F_PRICES, R.F_SURVEY) if re.search(r"2026 ?Q2|June quarter 2026",
                                                                                      text[f])]
    ok("SF18 no other June-quarter estimate of the tracker's lines ships",
       not other_q2 and set(poll) >= {"forecasts"} and "components" not in json.dumps(poll), str(other_q2))

    # ------------------------------------------------------------------------------------------------ the ladder
    R0, R1_, R2_ = routes["month-end"][1], routes["quarter average"][1], routes["two months for every industry"][1]
    ok("L1 the rungs land on their named figures: month-end 2.2, quarter average 3.3, two months 3.1, golden 2.8",
       (r1(R0), r1(R1_), r1(R2_), r1(G)) == (2.2, 3.3, 3.1, 2.8), "%.3f, %.3f, %.3f, %.4f" % (R0, R1_, R2_, G))
    cor = {k: v for k, (kind, v) in routes.items() if kind in ("C", "S")}
    ok("L2 every construction of the revaluation and both one-sigma moves sit in the golden's bin",
       all(in_bin(v) for v in cor.values()),
       "corridor %+.3f to %+.3f" % (min(cor.values()) - G, max(cor.values()) - G))
    riv = {k: v for k, (kind, v) in routes.items() if kind == "R"}
    near = min(riv, key=lambda k: abs(riv[k] - G))
    ok("L3 every rival lands outside the bin, at least 0.1 from the golden",
       all(not in_bin(v) and abs(v - G) >= 0.1 for v in riv.values()),
       "%d rivals; nearest %s %.3f (%+.3f)" % (len(riv), near, riv[near], riv[near] - G))
    ok("L4 the golden is bracketed: month-end below, the quarter's average above", R0 < G < R1_)
    long_t = [c for c in K.IND_CODES if min(B["ratio"][c][m] for m in ("2026-03", "2026-06")) > 1.0]
    short_t = [c for c in K.IND_CODES if c not in long_t]
    same_short = all(abs(C.stock_volume(B, c, m, "turnover") - C.stock_volume(B, c, m, "monthend")) < 1e-9
                     for c in short_t for m in (C.m_idx("2026-03"), C.m_idx("2026-06")))
    part = {}
    for n in range(len(long_t)):
        for sub in itertools.combinations(long_t, n):
            v = 0.0
            for c in K.IND_CODES:
                meth = "turnover" if c in sub else "monthend"
                v += (C.stock_volume(B, c, C.m_idx("2026-06"), meth) - C.stock_volume(B, c, C.m_idx("2026-03"), meth))
            part[sub] = C.call(B, v)["ann"]
    nearest = max(part, key=lambda k: part[k])
    ok("L5 a stock held under a month is the same under every window; no partial application lands in the bin",
       same_short and len(long_t) == 5 and len(part) == 31 and not any(in_bin(v) for v in part.values()),
       "%d subsets of %s; nearest %s %.3f" % (len(part), ",".join(long_t), "+".join(nearest), part[nearest]))
    grid = {}
    m1, m0 = C.m_idx("2026-06"), C.m_idx("2026-03")
    wins = {"month end": ("monthend", None), "quarter average": ("qavg", None), "two months": ("turnover", 2.0),
            "own ratio": ("turnover", None)}
    for ends in ("both", "June only", "March only"):
        for wn, (meth, T) in wins.items():
            v = 0.0
            for c in K.IND_CODES:
                e1 = meth if ends in ("both", "June only") else "monthend"
                e0 = meth if ends in ("both", "March only") else "monthend"
                v += C.stock_volume(B, c, m1, e1, T if e1 == "turnover" else None) - \
                    C.stock_volume(B, c, m0, e0, T if e0 == "turnover" else None)
            grid[(ends, wn)] = C.call(B, v)["ann"]
    bad = [k for k, v in grid.items() if k != ("both", "own ratio") and (in_bin(v) or abs(v - G) < 0.1)]
    ok("L6 the grid: every cell but each industry's own ratio at both ends is at least 0.1 from the golden",
       not bad and in_bin(grid[("both", "own ratio")]), "; ".join("%s/%s %.2f" % (a, b, v) for (a, b), v in
                                                                  grid.items()))

    # ------------------------------------------------------------------------------------------------ back-test
    t_rmse = bt["turnover"]["rmse"]
    rivals_bt = {k: v["rmse"] for k, v in bt.items() if v["kind"] == "R"}
    worst = min(rivals_bt, key=rivals_bt.get)
    ok("BT1 the revaluation over each industry's own ratio comes within $12m RMS of the Office's past estimates; "
       "every rival misses by three times that or more",
       t_rmse < 12.0 and all(v >= 3.0 * t_rmse for v in rivals_bt.values()),
       "turnover %.1f over %d quarters; best rival %s %.1f" % (t_rmse, bt["turnover"]["n"], worst, rivals_bt[worst]))
    ok("BT2 the March quarter 2026: month-end misses the Office by more than $100m, the revaluation by under $5m",
       abs(bt["month-end"]["last"]) > 100.0 and abs(bt["turnover"]["last"]) < 5.0,
       "%.1f and %.1f" % (bt["month-end"]["last"], bt["turnover"]["last"]))
    ok("BT3 no construction reproduces every quarter (the Office's estimate for businesses outside the survey)",
       all(v["max"] > 5.0 for v in bt.values()), "turnover's largest miss %.1f" % bt["turnover"]["max"])

    # ------------------------------------------------------------------------------------------------ Ask A
    gold_a = {k: (r1(v) if k in A.GRADED_A else r2(v)) for k, v in fa.items()}
    clear_a = {k: K.bin_clearance(v, 0.1 if k in A.GRADED_A else 0.01) for k, v in fa.items()}
    ok("A1 the revision record's five figures, each 0.12 of a bin or more from its edge",
       min(clear_a.values()) >= 0.12, "%s; clearances %s" % (gold_a, {k: round(v, 2) for k, v in clear_a.items()}))
    stops = A.ask_a_stops(B)
    missed = {}
    for name, st in stops.items():
        if name == "golden":
            continue
        missed[name] = sum((r1(st[k]) if k in A.GRADED_A else r2(st[k])) != gold_a[k] for k in gold_a)
    ok("A2 the stops: the 26 June download misses 5 of 5, mixed releases 5, current prices 5, the sign 4",
       [missed[k] for k in sorted(missed)] == [5, 5, 5, 4], str(missed))
    lat = B["latest"]["GDP"]
    revs4 = []
    for q in A.QUARTERS_A:
        adv = A.db_series(B, "%s-A" % q, "GDP")
        revs4.append((lat[q] / lat[A.prev_q(q)] - adv[q] / adv[A.prev_q(q)]) * 400.0)
    f4 = {q: v for q, v in zip(A.QUARTERS_A, revs4) if q in A.GRADED_A}
    f4["mean"], f4["mar"] = float(np.mean(revs4)), float(np.mean(np.abs(revs4)))
    m4 = sum((r1(f4[k]) if k in A.GRADED_A else r2(f4[k])) != gold_a[k] for k in gold_a)
    ok("A3 HA3 annualising by four instead of compounding (against the guide's compound annual rate) moves the two "
       "means and no quarter", m4 == 2 and all(r1(f4[q]) == gold_a[q] for q in A.GRADED_A), "%d moved" % m4)
    rr = []
    for q in A.QUARTERS_A:
        rnd = lambda x: K.round_half_up(x, 1)
        rr.append(rnd(rec[q][1]) - rnd(rec[q][0]))
    ok("A4 the guide measures revisions on unrounded growth rates (the rounded-headline reading moves the mean)",
       "Revisions are measured between growth rates calculated from the unrounded estimates, not between the rounded "
       "headlines." in guide, "rounded-first mean %.3f against %.3f" % (float(np.mean(rr)), fa["mean"]))
    last_rel = rel_ids[-1]
    upd = [q for q in A.QUARTERS_A if B["db"][(last_rel, "Chain", q, "GDP")] != lat[q]]
    pre = [q for q in NA.QUARTERS[NA.QI["2019Q1"]:NA.QI["2023Q2"]] if B["db"][(last_rel, "Chain", q, "GDP")] != lat[q]]
    ok("A5 A1: the download's last release is the 25 June third estimate; the 21 July update moved GDP in every "
       "quarter from June 2023 and in none before",
       last_rel == stale and B["releases"][last_rel][0] == "2026-06-25" and "2026-U" not in B["releases"] and
       len(upd) == 12 and not pre, "%s %s; %d quarters moved, %d before June 2023" % (
           last_rel, B["releases"][last_rel][0], len(upd), len(pre)))
    narr = r"not (?:yet )?in the (?:download|database)|predates|before the annual|stale|out of date|superseded"
    ok("A6 no shipped sentence says the download predates the update", not any(re.search(narr, t, re.I) for t in
                                                                                 text.values()))
    basis = {}
    with open(os.path.join(target, R.F_DB), encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            if row["price_basis"].startswith("Chain"):
                basis[row["release_id"]] = row["price_basis"]
    qlist = NA.QUARTERS
    rule_ok, n_rel, broke = True, 0, []
    for i, rid in enumerate(rel_ids):
        if i == 0 or basis[rid] != basis[rel_ids[i - 1]]:
            continue            # a March quarter's advance estimate re-expresses the series at the new year's prices
        prev = rel_ids[i - 1]
        kind = B["releases"][rid][1]
        if kind == "Annual update":
            first = "%dQ2" % (int(B["releases"][rid][0][:4]) - 3)
            window = {q for q in qlist if q >= first}
        else:
            own = rid[:6]
            k = qlist.index(own)
            span = {"Advance": 0, "Second": 1, "Third": 2}[kind]
            window = set(qlist[k - span:k + 1])
        n_rel += 1
        for q in qlist:
            if q in window:
                continue
            for sb_ in ("Chain", "Curre"):
                a_, b_ = B["db"].get((rid, sb_, q, "GDP")), B["db"].get((prev, sb_, q, "GDP"))
                if a_ is not None and b_ is not None and a_ != b_:
                    rule_ok = False
                    broke.append((rid, q))
    ok("A7 the guide's revision rules hold in the record: every release leaves the quarters outside its window as the "
       "release before it had them (the March quarters' advance estimates re-express the series)",
       rule_ok and n_rel >= 80, "%d releases checked; %s" % (n_rel, broke[:3]))

    # ------------------------------------------------------------------------------------------------ Ask B
    gold_b = {k: r1(v) for k, v in fb.items()}
    clear_b = {k: K.bin_clearance(v, 0.1) for k, v in fb.items()}
    ok("B1 the coal and LNG table's twelve figures, each 0.12 of a bin or more from its edge (exact sums of the "
       "shipped values)", len(fb) == 12 and min(clear_b.values()) >= 0.12, "min clearance %.3f" % min(clear_b.values()))
    sb = A.ask_b_stops(B)
    mv = {name: sum(r1(st[k]) != gold_b[k] for k in gold_b) for name, st in sb.items() if name != "golden"}
    ok("B2 the stops: the download as it stands misses 11 of 12, LNG moved alone 6, coal moved alone 5, two months 12",
       [mv[k] for k in sorted(mv)] == [11, 6, 5, 12], str(mv))
    fair = True
    for c in A.MARKETS:
        for d in ("32", "34"):
            truth = sum(W.SHIP_V[d][c][WD.MI[m]] for m in ("2025-04", "2025-05", "2025-06")) if c in W.SHIP_V[d] else 0
            fair &= abs(A.quarter_2025(B, c, d) - truth) < 0.01
            t26 = sum(W.SHIP_V[d][c][WD.MI[m]] for m in ("2026-04", "2026-05", "2026-06")) if c in W.SHIP_V[d] else 0
            fair &= abs(A.quarter_2026(B, c, d) - t26) < 0.1
    ok("B3 fair: moving each 2025 month back one reproduces the shipments exactly; 2026 is on shipments", fair)
    lag = A.lag_evidence(B)
    g25 = [(abs(lag[m][d][0]), abs(lag[m][d][1])) for m in K.month_list("2025-02", "2025-11") for d in ("32", "34")]
    g26 = [(abs(lag[m][d][0]), abs(lag[m][d][1])) for m in K.month_list("2026-01", "2026-05") for d in ("32", "34")]
    ok("B4 the referee: in 2025 TBS's coal and LNG tonnes match the ports' loadings a month earlier, from January 2026 "
       "the same month",
       max(b for a, b in g25) < 0.02 and min(a for a, b in g25) > 0.0 and np.mean([a for a, b in g25]) > 0.05 and
       max(a for a, b in g26) < 0.02,
       "2025: same month %.3f, month before %.3f; 2026: same month %.3f (mean absolute gaps)" % (
           np.mean([a for a, b in g25]), np.mean([b for a, b in g25]), np.mean([a for a, b in g26])))
    dec_t = {d: A.tbs_tonnes(B, "2025-12", d) for d in ("32", "34")}
    dec_p = {d: A.port_tonnes(B, "2025-11", cg) + A.port_tonnes(B, "2025-12", cg)
             for d, cg in (("32", "Coal"), ("34", "LNG"))}
    ok("B5 December 2025 carries November's and December's cargoes",
       all(abs(dec_t[d] / dec_p[d] - 1) < 0.02 for d in dec_t),
       "; ".join("%s %.3f" % (d, dec_t[d] / dec_p[d]) for d in dec_t))
    with open(os.path.join(target, R.F_BULK), encoding="utf-8") as fh:
        bh = fh.readline().strip().split(",")
    marker = r"periodic|lodgement month|month (?:they|it) (?:was|were) lodged|month of shipment|shipment month|" \
             r"by (?:their|its) shipments|sailed"
    ok("B6 no marker or narration: the bulk download's columns are the same for every division and no shipped text "
       "describes how periodic entries are dated", bh == R.BULK_COLS and
       not any(re.search(marker, t, re.I) for t in text.values()), ", ".join(bh))
    kt = {c: B["bulletin"]["energy"][c] for c in A.MARKETS}
    ok("B7 HB1 the download is in thousands of dollars, the bulletin in millions: unread, every figure moves",
       all(abs(A.bulk_value(B, "2026-05", c, "32") * 1000 - B["bulk"][("2026-05", "Exports", c, "32")][0]) < 1
           for c in A.MARKETS) and "Values are in millions of dollars." in bull)
    sep_x = all(B["bulletin"]["exp"]["3"][m] == round(float(W.SHIP_VX["3"][WD.MI[m]]), 1) for m in R.BULLETIN_MONTHS)
    q1_x = sum(B["portal_xsa"]["3"][m] == round(float(W.SHIP_VX["3"][WD.MI[m]]), 1) for m in ("2026-01", "2026-02"))
    ok("B8 the separation line: the call's goods months (January to June 2026) carry no lodgement shift",
       sep_x and q1_x == 2 and abs(gq1["x"] - B["base"]["X_G"]) < 1.0)

    # ------------------------------------------------------------------------------------------------ separation
    Bsep = {k: v for k, v in B.items() if k not in ("bulk", "ports", "countries", "releases")}
    Bsep["db"] = {k: v for k, v in B["db"].items() if k[3] == "CIPI_NF"}
    ok("SEP1 the call never reads an ask's rows: the bulk download, the ports, the database's GDP rows deleted",
       C.golden(Bsep)["ann"] == G)
    Bpert = dict(B)
    Bpert["stocks"] = {c: {m: v * 1.05 for m, v in d.items()} for c, d in B["stocks"].items()}
    Bpert["p25"] = {c: v * 1.03 for c, v in B["p25"].items()}
    ok("SEP2 the asks never read the call's rows: the stocks and input prices perturbed, both asks unchanged",
       A.ask_a(Bpert)[0] == fa and A.ask_b(Bpert) == fb)

    # ------------------------------------------------------------------------------------------------ realism
    import subprocess
    repo = os.path.dirname(TASK)
    with open(os.path.join(repo, ".claude", "skills", "fingerprint", "cards", "task154.json"), encoding="utf-8") as fh:
        planned = json.load(fh)["world"]["people"]
    pr = subprocess.run(["python3", os.path.join(repo, ".claude", "skills", "fingerprint", "guard.py"), "people",
                         "task154"], capture_output=True, text=True, cwd=repo)
    gtext = pack_text(golden)
    alltext = " ".join(text.values()) + " " + " ".join(gtext.values())
    people = [p for p in K.PEOPLE.values() if p in alltext]
    ok("RS1 every person named in the pack and the goldens is on the card, drawn by guard.py names (en_GB, seed 154, "
       "20), and guard.py people clears them", people and all(p in planned for p in people) and pr.returncode == 0 and
       "BLOCK" not in pr.stdout, "%s; guard: %s" % (", ".join(people), pr.stdout.strip().splitlines()[0][:80]))
    tab = B["bulletin"]
    vals1 = [v for key in ("exp", "imp", "stocks", "sales") for d in tab[key].values() for v in d.values()]
    zero_dec = sum(1 for v in vals1 if abs(v * 10 - round(v * 10)) < 1e-6 and round(v * 10) % 10 == 0) / len(vals1)
    totals = [sum(tab[key][x][m] for x in tab[key]) for key in ("exp", "imp", "stocks", "sales")
              for m in R.BULLETIN_MONTHS]
    base_round = sum(1 for v in B["base"].values() if v % 100 == 0)
    ok("RS2 no round figures: the bulletin's values end in .0 at a natural rate, no total is a multiple of $10m, no "
       "two industries or sections share a level",
       zero_dec < 0.2 and not any(abs(t / 10 - round(t / 10)) < 1e-6 for t in totals) and base_round <= 1 and
       len({tab["stocks"][c]["2026-03"] for c in K.IND_CODES}) == 8,
       "%.0f%% end in .0; %d base levels in whole hundreds" % (100 * zero_dec, base_round))
    ports = B["ports"]
    vessels = [v for (_, _, cg, _, v) in ports if cg == "Coal"]
    ok("RS3 no constant stream: the ports' monthly loadings and vessel counts vary", len(set(vessels)) > 10 and
       len({t for (_, _, _, t, _) in ports}) == len(ports))
    h1 = []
    for folder, names in ((target, files), (golden, sorted(os.listdir(golden)))):
        for f in names:
            p = os.path.join(folder, f)
            if f.endswith(".xlsx"):
                z = zipfile.ZipFile(p)
                core = z.read("docProps/core.xml").decode()
                app = z.read("docProps/app.xml").decode()
                stamps = re.findall(r"<dcterms:(?:created|modified)[^>]*>([^<]+)<", core)
                who = re.findall(r"<(?:dc:creator|cp:lastModifiedBy)>([^<]*)<", core)
                if "openpyxl" in core + app or "Microsoft Excel" not in app or \
                        any(not ("2026-04-01" <= s[:10] <= "2026-07-29") for s in stamps) or \
                        any(w not in (K.OFFICE, K.PEOPLE["analyst"], K.PEOPLE["economist"]) for w in who):
                    h1.append(f)
            elif f.endswith(".pdf"):
                import pypdf
                md = pypdf.PdfReader(p).metadata
                raw = open(p, "rb").read()
                st = [md.get("/CreationDate", ""), md.get("/ModDate", "")]
                if b"ReportLab" in raw or b"reportlab" in raw or md.get("/Producer") or md.get("/Creator") or \
                        any(not ("D:20260401" <= s[:10] <= "D:20260729") for s in st):
                    h1.append(f)
    ok("H1 no writer's name or build-clock stamp in any workbook or PDF, pack or goldens", not h1, str(h1))
    late = []
    for f in files:
        for d in re.findall(r"\b(20\d\d)-(\d\d)(?:-(\d\d))?\b", text[f]):
            s = "%s-%s-%s" % (d[0], d[1], d[2] or "01")
            if s > "2026-07-29":
                late.append((f, s))
    mnames = "|".join(K.MONTH_NAMES[1:])
    ahead = set()
    for f in files:
        for d, mo, y in re.findall(r"\b(\d{1,2}) (%s) (20\d\d)\b" % mnames, text[f]):
            if "%s-%02d-%02d" % (y, K.MONTH_NAMES.index(mo), int(d)) > "2026-07-29":
                ahead.add(f)
    mt = {os.path.getmtime(os.path.join(target, f)) for f in files}
    ok("H16 no date in the pack runs past the desk's folder of 29 July but the release calendars' (the guide, the "
       "notice, the poll's target); every file stamped at the export",
       not late and ahead <= {R.F_GUIDE, R.F_NOTICE, R.F_POLL} and len(mt) == 1, "%s; ahead %s" % (late[:4],
                                                                                              sorted(ahead)))
    invented = [K.OFFICE, K.TBS, K.TBS_DEPT, K.BANK, K.DESK, K.WIRE, "Harbours Board", "Industry Federation"] + \
        sorted({p for (_, p, _, _, _) in ports})
    hits = []
    for name in invented:
        for real in REAL:
            for a in distinctive(name):
                for b in distinctive(real) or [real.lower()]:
                    if edit_distance(a, b) <= 2 or (len(b) > 3 and (a in b or b in a)):
                        hits.append((name, real))
    ok("H21 no invented organisation or port sits within two letters of a real one or contains its distinctive word",
       not hits, "%d names against %d; %s" % (len(invented), len(REAL), hits[:4]))

    with open(os.path.join(target, R.F_INDEX), encoding="utf-8") as fh:
        listed = [row["file"] for row in csv.DictReader(fh)]
    ok("H9 the desk's data index lists every shipped file but itself, once each",
       sorted(listed) == sorted(f for f in files if f != R.F_INDEX) and len(listed) == len(set(listed)),
       "%d listed" % len(listed))
    tree = [os.path.join(dp, fn) for dp, dn, fns in os.walk(root) for fn in fns]
    ok("H11 one submission.md and one prompt.txt in the task tree, no prompt.md, no markdown in the pack",
       sum(1 for x in tree if x.endswith(os.sep + "submission.md")) == 1 and
       sum(1 for x in tree if x.endswith(os.sep + "prompt.txt")) == 1 and
       not any(x.endswith("prompt.md") for x in tree) and not any(f.endswith(".md") for f in files))
    built = r"constructed|fictional|fictitious|invented|this exercise|synthetic|generated (?:for|by)|made up"
    ok("H24 nothing under target/ says the pack was built", not any(re.search(built, t, re.I) for t in text.values()))

    # ------------------------------------------------------------------------------------------------ one file each
    facts = {"Stocks are at cost.": R.F_BULLETIN,
             "The stock-to-sales ratio is stocks at the end of the month divided by the month's cost of sales":
                 R.F_BULLETIN,
             "The advance estimate revises no earlier quarter": R.F_GUIDE,
             "((1 + q/100) to the fourth power, less 1) times 100": R.F_GUIDE,
             "The advance estimate will incorporate the June results.": R.F_NOTICE,
             "producer price indexes, inputs to each industry": R.F_GUIDE,
             "Revisions are measured between growth rates calculated from the unrounded estimates": R.F_GUIDE,
             "From the January 2026 month the statistics are compiled on TBS's new processing system.": R.F_BULLETIN,
             "System of National Accounts 2008": R.F_GUIDE}
    where = {k: [f for f in files if k in text[f]] for k in facts}
    ok("ONE every load-bearing sentence sits in exactly one file, the one it belongs to",
       all(where[k] == [v] for k, v in facts.items()), str({k[:30]: v for k, v in where.items() if v != [facts[k]]}))

    # ------------------------------------------------------------------------------------------------ distractors
    tmp = tempfile.mkdtemp(dir=root, suffix=".dis")
    try:
        for f in files:
            if f not in meta["distractor_files"]:
                shutil.copy2(os.path.join(target, f), os.path.join(tmp, f))
        Bd = R.bundle(tmp, W)
        ok("DIS with both distractors deleted the call and both asks are unchanged",
           C.golden(Bd)["ann"] == G and A.ask_a(Bd)[0] == fa and A.ask_b(Bd) == fb)
    finally:
        shutil.rmtree(tmp)

    # ------------------------------------------------------------------------------------------------ goldens
    pdf = gtext["gdp_preview_2026q2.pdf"]
    head = "June quarter GDP: we expect %.1f%% annualised" % r1(G)
    first = "to show real GDP growth of %.1f%% at an annualised rate in the June quarter" % r1(G)
    ok("FIG1 the preview opens with the call: the headline and the first sentence",
       pdf.find(head) != -1 and pdf.find(first) != -1 and pdf.find(head) < pdf.find(first) < 400,
       "headline at %d, sentence at %d" % (pdf.find(head), pdf.find(first)))
    gw = openpyxl.load_workbook(os.path.join(golden, "gdp_preview_2026q2_tables.xlsx"), data_only=True)
    rr_ = gw["Revision record"]
    graded = {}
    for r_ in range(1, rr_.max_row + 1):
        lab = rr_.cell(row=r_, column=1).value
        if isinstance(lab, str) and (lab.startswith("Revision, ") or lab.startswith("Mean ")):
            graded[lab] = rr_.cell(row=r_, column=4).value
    want = [figs["ask_a"][k] for k in ("2025Q3", "2025Q4", "2026Q1", "mean", "mar")]
    cl = gw["Coal and LNG"]
    cells = {}
    for r_ in range(5, 5 + len(A.MARKETS) + 1):
        cells[cl.cell(row=r_, column=1).value] = (cl.cell(row=r_, column=4).value, cl.cell(row=r_, column=7).value)
    wantb = {A.MARKET_NAMES[c]: (figs["ask_b"]["%s_32" % c], figs["ask_b"]["%s_34" % c]) for c in A.MARKETS}
    wantb["Five markets"] = (figs["ask_b"]["total_32"], figs["ask_b"]["total_34"])
    ok("FIG2 the workbook's graded cells equal figures.json: the call, the revision record and the coal and LNG table",
       gw["Call"]["B3"].value == figs["call"]["annualised_pct"] == r1(G) and list(graded.values()) == want and
       cells == wantb, "%s; %s" % (list(graded.values()), cells.get("Five markets")))
    ok("FIG3 the call's components agree everywhere: contributions sum to the quarter's growth",
       abs(sum(g["contrib"].values()) - g["q"]) < 1e-9 and figs["call"]["quarterly_pct"] == r2(g["q"]))

    # ------------------------------------------------------------------------------------------------ pair arithmetic
    crack_a = sum((r1(stops["S1 the download's 25 June release as the latest"][k]) if k in A.GRADED_A else
                   r2(stops["S1 the download's 25 June release as the latest"][k])) == gold_a[k] for k in gold_a)
    crack_b = sum(r1(sb["S1 the download as it stands"][k]) == gold_b[k] for k in gold_b)
    n_ask = len(gold_a) + len(gold_b)
    lc = (crack_a + crack_b) / n_ask
    cc = 38 + 7 + 55 * lc
    ss = 5 + 7 + 55 * lc
    pair = (cc + ss) / 2
    pair_a1 = (38 + 7 + 55 * (len(gold_a) + crack_b) / n_ask + ss) / 2
    ok("PAIR the cracker keeps %d of %d ask figures: mirror leakage under 35 per cent, pair ceiling under 40, and "
       "under 40 with A1 cracked too" % (crack_a + crack_b, n_ask),
       100 * lc < 35 and pair <= 40 and pair_a1 <= 40,
       "leakage %.1f%%; pair %.1f; with A1 cracked %.1f" % (100 * lc, pair, pair_a1))

    # ------------------------------------------------------------------------------------------------ report
    fails = [r_ for r_ in ok.rows if not r_[1]]
    for label, c, d in ok.rows:
        print("%s  %s%s" % ("PASS" if c else "FAIL", label, (": " + d) if d else ""))
    print("checks: %d run, %d failed" % (len(ok.rows), len(fails)))
    if fails:
        raise SystemExit("checks failed: " + "; ".join(f[0] for f in fails))
    return ok.rows
