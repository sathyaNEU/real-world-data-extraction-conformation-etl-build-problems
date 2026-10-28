"""The goldens (the 2027 GPS order workbook and the Category V forecast memo), generator/figures.json, submission.md and
metadata.json, all computed from the shipped files through calc.py."""
from __future__ import annotations

import datetime as dt
import json
import os

import numpy as np

import calc
import common as C

MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
          "November", "December"]
JAILS = {"D": "Darrow County Jail", "K": "Keswick County Jail"}
GOLDEN_FORM = "victim 12 months"


def compute(target):
    s = calc.Solution(target)
    o = dict(s=s)
    o["forms"] = s.all_forms()
    o["golden"] = o["forms"][GOLDEN_FORM][0]
    o["default"] = o["forms"]["own attributes"][0]
    f26 = s.f26
    o["caseload"] = int(len(f26))
    o["victims_2026"] = int(f26["v12"].sum())
    f25 = s.frames[2025]
    o["victims_2025"] = int(f25["v12"].sum())
    tr = s.train
    o["rate_victims"] = float(tr.loc[tr["v12"] == 1, "y"].mean())
    o["rate_others"] = float(tr.loc[tr["v12"] == 0, "y"].mean())
    o["n_victim_years"] = int((tr["v12"] == 1).sum())
    h = s.d["hpd"]
    hi = h[(h["victim_type"] == "I") & (h["supplement_no"] == "0")].drop_duplicates(["incident_number", "victim_seq"])
    yr = hi["occurred_date"].str[6:10].astype(int).value_counts()
    o["city_victims"] = {2025: int(yr[2025]), 2026: int(yr[2026])}
    ch = yr[2026] / yr[2025] - 1
    o["city_change_word"] = ("fell %.1f per cent" % (-100 * ch)) if ch < -0.005 else ("rose %.1f per cent" % (100 * ch) if ch > 0.005 else "was flat")
    o["gps"], o["gps_detail"] = calc.gps_days(s)
    o["zones"] = calc.zones(s)
    nb = s.panel.nb
    o["schedb_bookings"] = int(nb.loc[nb["schedb"] & (nb["date"] >= C.RECORDS_START) & (nb["date"] <= C.RECORDS_END)].shape[0])
    return o


def figures(o):
    f = {"call_forecast_2027": C.nearest(o["golden"], C.ROUND_CALL)}
    f["golden_unrounded"] = round(o["golden"], 1)
    f["default_unrounded"] = round(o["default"], 1)
    f["default_printed"] = C.nearest(o["default"], C.ROUND_CALL)
    f["cc_caseload_2026"] = o["caseload"]
    f["cc_victims_on_caseload_2026"] = o["victims_2026"]
    f["cc_victims_on_caseload_2025"] = o["victims_2025"]
    # the victim counts move by a few with the match key and the date field, so the write-up states them to the nearest 10
    f["cc_victims_on_caseload_2026_about"] = C.nearest(o["victims_2026"], 10)
    f["cc_victims_on_caseload_2025_about"] = C.nearest(o["victims_2025"], 10)
    f["cc_rate_victims_pct"] = round(100 * o["rate_victims"], 1)
    f["cc_rate_others_pct"] = round(100 * o["rate_others"], 1)
    f["step_schedb_bookings_2018_2026"] = o["schedb_bookings"]
    f["years_train"] = [2018, 2025]
    f["year_end_origin"] = 2026
    f["round_call_to"] = C.ROUND_CALL
    for m in range(12):
        f["ask_gps_days_2026_%02d" % (m + 1)] = int(o["gps"][m])
    f["ask_gps_days_2026_total"] = int(sum(o["gps"]))
    for m in range(12):
        f["ask_zones_2026_%02d" % (m + 1)] = int(o["zones"][m])
    f["ask_zones_2026_total"] = int(sum(o["zones"]))
    return f


# ------------------------------------------------------------------ goldens
def write_workbook(path, o, f):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    import postprocess as PP
    wb = Workbook()
    ws = wb.active
    ws.title = "2027 reservation"
    bold = Font(bold=True)
    hdr = PatternFill("solid", fgColor="D9E1F2")
    ws["A1"] = "Region 4: 2027 Belvantis order, enrolments reserved under Directive 26-07"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Prepared 8 January 2027 from the records to 31 December 2026"
    rows = [("People on the Region 4 caseload at the end of 31 December 2026", o["caseload"]),
            ("Forecast: of those, booked on a new Category V charge in 2027 while still supervised", f["call_forecast_2027"]),
            ("Enrolments to reserve for 2027", f["call_forecast_2027"])]
    for k, (a, b) in enumerate(rows, 4):
        ws.cell(row=k, column=1, value=a)
        c = ws.cell(row=k, column=2, value=b)
        c.number_format = "#,##0"
    ws.cell(row=5, column=1).font = bold
    ws.cell(row=5, column=2).font = bold
    ws["A8"] = ("Forecast to the nearest 100; see the forecast memo for the model. Exclusion zones for 2027 are not "
                "forecast here.")
    ws.column_dimensions["A"].width = 84
    ws.column_dimensions["B"].width = 12
    ws2 = wb.create_sheet("Enrolment days 2026")
    ws2["A1"] = "Enrolment days owed to Belvantis for 2026, by month (Schedule C)"
    ws2["A1"].font = Font(bold=True, size=12)
    ws2["A3"] = "Month"
    ws2["B3"] = "Enrolment days"
    for c in ("A3", "B3"):
        ws2[c].font = bold
        ws2[c].fill = hdr
    for m in range(12):
        ws2.cell(row=4 + m, column=1, value="%s 2026" % MONTHS[m])
        c = ws2.cell(row=4 + m, column=2, value=f["ask_gps_days_2026_%02d" % (m + 1)])
        c.number_format = "#,##0"
    ws2.cell(row=16, column=1, value="Total 2026").font = bold
    c = ws2.cell(row=16, column=2, value="=SUM(B4:B15)")
    c.number_format = "#,##0"
    c.font = bold
    ws2["A18"] = ("Days on which a client held a device on the Region 4 account, less suspended days, days in either county "
                  "jail and days after Region 4's supervision ended; voided events disregarded, and a client Belvantis holds "
                  "under both a legacy offender number and a CMS id counted once.")
    ws2.column_dimensions["A"].width = 24
    ws2.column_dimensions["B"].width = 16
    ws3 = wb.create_sheet("Exclusion zones 2026")
    ws3["A1"] = "Exclusion zones Directive 26-07 section 4 would have set in 2026, by month"
    ws3["A1"].font = Font(bold=True, size=12)
    ws3["A3"], ws3["B3"] = "Month", "Zones set"
    for c in ("A3", "B3"):
        ws3[c].font = bold
        ws3[c].fill = hdr
    for m in range(12):
        ws3.cell(row=4 + m, column=1, value="%s 2026" % MONTHS[m])
        c = ws3.cell(row=4 + m, column=2, value=f["ask_zones_2026_%02d" % (m + 1)])
        c.number_format = "#,##0"
    ws3.cell(row=16, column=1, value="Total 2026").font = bold
    c = ws3.cell(row=16, column=2, value="=SUM(B4:B15)")
    c.number_format = "#,##0"
    c.font = bold
    ws3["A18"] = ("Halverton Police Department arrests of Region 4 supervisees on Category V charges, at the first booking on "
                  "each incident; one zone per supervisee for each Halverton residence of the incident's victims, a "
                  "residence counted once however the report and its supplements write it.")
    for col, wv in (("A", 24), ("B", 14)):
        ws3.column_dimensions[col].width = wv
    ws4 = wb.create_sheet("Notes")
    notes = [("As of", "Records to 31 December 2026; prepared 8 January 2027."),
             ("Caseload", "People with an open Region 4 episode at the end of 31 December 2026, counted once (CMS export)."),
             ("Category V", "Directive 26-07, Appendix; former Title 21 and Title 8 sections read through the Revisor's 2024 table."),
             ("Forecast", "New-arrest bookings in either county jail in 2027 while still supervised, people booked at least once."),
             ("Enrolment day", "Belvantis contract MDOC-2025-117, Schedule C: a day a client holds a device on the Region 4 "
                               "account, less suspended days, custody days and days after Region 4's supervision ended."),
             ("Exclusion zone", "Directive 26-07, section 4: one per Halverton residence of the incident's victims, per "
                                "supervisee and incident, set in the month of the first booking."),
             ("Prepared by", "Zachary Wallace, Research and Planning, Region 4.")]
    ws4["A1"] = "Notes"
    ws4["A1"].font = Font(bold=True, size=12)
    for k, (a, b) in enumerate(notes, 3):
        ws4.cell(row=k, column=1, value=a).font = bold
        ws4.cell(row=k, column=2, value=b)
    ws4.column_dimensions["A"].width = 16
    ws4.column_dimensions["B"].width = 110
    for w_ in (ws, ws2, ws3):
        w_.sheet_view.showGridLines = False
    ws2.freeze_panes = "A4"
    ws3.freeze_panes = "A4"
    wb.save(path)
    PP.normalize_ooxml(path, "Zachary Wallace", "Zachary Wallace", "2027-01-07T15:22:00Z", "2027-01-08T16:05:00Z",
                       excel=True, company=C.DEPT)


def write_memo(path, o, f):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor, Cm
    import postprocess as PP
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.3)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)
    fmt = lambda x: "{:,}".format(x)
    t = doc.add_paragraph()
    r = t.add_run("MEMORANDUM")
    r.bold = True
    r.font.size = Pt(14)
    for k, v in (("To", "Region 4 Executive Team"), ("From", "Zachary Wallace, Research and Planning"),
                 ("Date", "8 January 2027"), ("Subject", "2027 Belvantis order: enrolments for Directive 26-07")):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run("%s:\t" % k).bold = True
        p.add_run(v)
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("Reserve %s enrolments: about %s of the %s people on the caseload at 31 December 2026 will be booked on a "
                  "new Category V charge in 2027" % (fmt(f["call_forecast_2027"]), fmt(f["call_forecast_2027"]), fmt(o["caseload"])))
    r.bold = True
    r.font.size = Pt(12)
    doc.add_paragraph(
        "From 1 January Directive 26-07 puts every supervisee booked on a new Category V charge on GPS. Of the %s people "
        "on our caseload at the end of 31 December 2026, %s are expected to be booked into the Darrow County or Keswick "
        "County jail on such a charge at least once in 2027 while still under Region 4 supervision, to the nearest 100. "
        "That is the number of enrolments the order should reserve at the reserved rate." % (fmt(o["caseload"]), fmt(f["call_forecast_2027"])))
    doc.add_paragraph(
        "The figure is higher than our usual year-end risk model gives on its own (%s). The difference is violence "
        "against the people we supervise. Matching the Halverton Police Department's violent-incident victims to the "
        "caseload by name and date of birth, about %s people on the 2026 year-end caseload were recorded as victims in "
        "2026, against about %s on the 2025 year-end caseload, most of them from September to November. Across the 2018 to 2025 "
        "year ends, supervisees recorded as victims in the prior twelve months were booked on a Category V charge the "
        "following year at %.1f per cent, against %.1f per cent for everyone else, and the model without the police "
        "records under-predicts them in every one of those years. None of this shows in our own records: the city's "
        "count of victims %s in 2026, and the caseload's ages, offences, scores and terms look like any other year's."
        % (fmt(f["default_printed"]), fmt(f["cc_victims_on_caseload_2026_about"]), fmt(f["cc_victims_on_caseload_2025_about"]),
           f["cc_rate_victims_pct"], f["cc_rate_others_pct"], o["city_change_word"]))
    doc.add_paragraph(
        "Method. Both jails' booking notices were put on one clock (Keswick's are in UTC) and one record per booking, new "
        "arrests only, with charges under former Title 21 and Title 8 sections read through the Revisor's 2024 conversion "
        "table.\u00b9 Each year end's caseload comes from the case-management export, each person counted once across "
        "concurrent cases and county transfers. The model is a logistic regression of a Category V new-arrest booking in "
        "the following calendar year while still supervised, fitted on the 2018 to 2025 year ends, on age, sex, prior "
        "felonies, lead offence, supervision type, county, Halverton residence, months to expiration, first year of "
        "supervision, the latest MRAS total, the prior year's new-arrest and Category V bookings and a police victim "
        "record in the prior twelve months. The 2026 caseload is scored with its own 2026 records and the probabilities "
        "summed.")
    rows = [("Caseload at the end of 31 December 2026", fmt(o["caseload"])),
            ("Recorded by HPD as victims of violence in 2026 (nearest 10)", "about " + fmt(f["cc_victims_on_caseload_2026_about"])),
            ("The same, 2025 year-end caseload, in 2025 (nearest 10)", "about " + fmt(f["cc_victims_on_caseload_2025_about"])),
            ("Category V booking rate the next year, victims, 2018 to 2025 year ends", "%.1f%%" % f["cc_rate_victims_pct"]),
            ("Category V booking rate the next year, everyone else", "%.1f%%" % f["cc_rate_others_pct"]),
            ("Forecast for 2027, with the police records", fmt(f["call_forecast_2027"])),
            ("Forecast for 2027, own records only", fmt(f["default_printed"]))]
    tb = doc.add_table(rows=len(rows) + 1, cols=2)
    tb.style = "Table Grid"
    tb.cell(0, 0).text = "Forecast basis"
    tb.cell(0, 1).text = "Figure"
    for k, (a, b) in enumerate(rows, 1):
        tb.cell(k, 0).text = a
        tb.cell(k, 1).text = b
        tb.cell(k, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for c in tb.rows[0].cells:
        for run in c.paragraphs[0].runs:
            run.bold = True
    src = doc.add_paragraph("Sources: CMS episode export; Darrow and Keswick booking notices; HPD violent-incident extract; "
                            "MRAS assessments. Forecasts to the nearest 100.")
    src.runs[0].font.size = Pt(8.5)
    doc.add_paragraph(
        "Also in the order workbook: the 2026 enrolment days we owe Belvantis by month (%s for the year, Schedule C) and "
        "the exclusion zones section 4 of the directive would have set each month of 2026 for the Halverton police's "
        "Category V arrests of our supervisees (%s in all)." % (fmt(f["ask_gps_days_2026_total"]), fmt(f["ask_zones_2026_total"])))
    n = doc.add_paragraph("\u00b9 Directive 26-07, section 2. Keswick's charge table kept the former sections until its update "
                          "of 1 April 2025.")
    n.runs[0].font.size = Pt(8.5)
    doc.save(path)
    PP.normalize_docx(path, "Zachary Wallace", "Zachary Wallace", "2027-01-07T16:40:00Z", "2027-01-08T15:58:00Z", 4, 96, "t179",
                      company=C.DEPT)


# ------------------------------------------------------------------ write-up
def submission(o, f):
    g = f
    z = o["zones"]
    L = []
    L.append("## Tags")
    L.append("")
    L.append("**Domain:** Criminology & Justice (recidivism risk assessment).")
    L.append("**Analytical objective:** Forecasting & Predictive Modeling.")
    L.append("")
    L.append("## 1. Final Recommendation")
    L.append("")
    L.append("**Reserve %s GPS enrolments in Region 4's 2027 Belvantis order: %s people on the 31 December 2026 caseload will "
             "be booked on a new Category V charge in 2027 while still supervised.** Not the %s a risk model on the "
             "caseload's own records gives, which misses the roughly %s supervisees the Halverton police recorded as "
             "victims of violence in 2026."
             % ("{:,}".format(g["call_forecast_2027"]), "{:,}".format(g["call_forecast_2027"]),
                "{:,}".format(g["default_printed"]), "{:,}".format(g["cc_victims_on_caseload_2026_about"])))
    L.append("")
    L.append("## 2. Critical Components")
    L.append("")
    L.append("1. The caseload at the end of 31 December 2026 is **%s** people." % "{:,}".format(g["cc_caseload_2026"]))
    L.append("2. About **%s** of them were recorded by the Halverton police as victims of a violent "
             "incident that occurred in 2026, against about **%s** on the 2025 year-end caseload for incidents in 2025."
             % ("{:,}".format(g["cc_victims_on_caseload_2026_about"]), "{:,}".format(g["cc_victims_on_caseload_2025_about"])))
    L.append("3. On the 2018 to 2025 year ends, supervisees with a police victim record in the prior twelve months were "
             "booked on a Category V charge the next year at **%.1f%%**, against **%.1f%%** for the rest."
             % (g["cc_rate_victims_pct"], g["cc_rate_others_pct"]))
    L.append("")
    L.append("## 3. Step-by-Step Solution")
    L.append("")
    steps = [
        "Read %s (one booking per booking_no) and %s (UTC converted to Merrow time), kept new arrests per Directive "
        "26-07 section 2, read former sections through %s and flagged Category V charges by the directive's Appendix: "
        "%s Category V new-arrest bookings of supervisees from 2018 to 2026."
        % (C.F_DARROW, C.F_KESWICK, C.F_CODE, "{:,}".format(g["step_schedb_bookings_2018_2026"])),
        "Built each year end's caseload from %s, an episode open from its start date through its close date and each "
        "person counted once: %s people at the end of 2026." % (C.F_EPIS, "{:,}".format(g["cc_caseload_2026"])),
        "Matched the individual victims in %s to the caseload on surname, the first word of the given name and date of "
        "birth, the names in capitals with the suffixes JR, SR, II, III and IV and every character but a letter dropped, "
        "each victimization dated by the incident's occurred date: about %s of the 2026 year-end caseload victims in "
        "2026, about %s of the 2025 year-end caseload in 2025."
        % (C.F_HPD, "{:,}".format(g["cc_victims_on_caseload_2026_about"]), "{:,}".format(g["cc_victims_on_caseload_2025_about"])),
        "Fitted a logistic model of a Category V new-arrest booking in the next calendar year while supervised on the "
        "2018 to 2025 year ends (age, sex, prior felonies, lead offence family, supervision type, county, city residence, "
        "months to expiration, first year of supervision, latest MRAS total with legacy numbers through %s, prior-year "
        "new-arrest and Category V bookings, and a police victim record in the prior twelve months), where victims "
        "book at %.1f%% against %.1f%%." % (C.F_XWALK, g["cc_rate_victims_pct"], g["cc_rate_others_pct"]),
        "Scored the 2026 year-end caseload and summed the probabilities: %.1f, %s to the nearest 100 (the same model "
        "without the victim records gives %.1f)." % (g["golden_unrounded"], "{:,}".format(g["call_forecast_2027"]),
                                                     g["default_unrounded"]),
        "Rebuilt each client's held days in %s from effective dates with voided events dropped, a legacy offender "
        "number and its case-management id per %s read as one client, removed suspended days, every day the client "
        "spent any part of in either jail per both booking files and every day after Region 4's supervision ended per "
        "%s, a transfer between Darrow and Keswick counties continuing it (Schedule C, %s): %s enrolment days in 2026."
        % (C.F_GPS, C.F_XWALK, C.F_EPIS, C.F_SCHED, "{:,}".format(g["ask_gps_days_2026_total"])),
        "For each supervisee's Halverton police Category V arrest, once per supervisee and incident at its first booking, "
        "counted the distinct Halverton residences of the incident's individual victims in %s across the original "
        "report and its supplements (Directive 26-07 section 4, unit marks and street types read alike) in the month of "
        "that booking: %s zones in 2026." % (C.F_HPD, "{:,}".format(g["ask_zones_2026_total"])),
        "Reserve %s enrolments for 2027." % "{:,}".format(g["call_forecast_2027"]),
    ]
    for k, s_ in enumerate(steps, 1):
        L.append("%d. %s" % (k, s_))
    L.append("")
    L.append("## 4. Deliverable Answers")
    L.append("")
    L.append("### %s" % C.DELIV_DOCX)
    L.append("")
    L.append("1. %s people booked on a new Category V charge in 2027 while still supervised, to the nearest 100." %
             "{:,}".format(g["call_forecast_2027"]))
    L.append("")
    L.append("### %s" % C.DELIV_XLSX)
    L.append("")
    L.append("1. Forecast: %s." % "{:,}".format(g["call_forecast_2027"]))
    L.append("2. Enrolment days owed to Belvantis for 2026, whole days:")
    for m in range(12):
        L.append("   - %s: %s" % (MONTHS[m], "{:,}".format(g["ask_gps_days_2026_%02d" % (m + 1)])))
    L.append("3. Exclusion zones the directive would have set in 2026 for the Halverton police's Category V arrests of "
             "supervisees, whole zones:")
    for m in range(12):
        L.append("   - %s: %s" % (MONTHS[m], "{:,}".format(g["ask_zones_2026_%02d" % (m + 1)])))
    return "\n".join(L) + "\n"


def metadata(n):
    src = {}
    made = {C.F_EPIS: "case-management export of every supervision episode open at any time in 2018 to 2026",
            C.F_XWALK: "the June 2022 conversion of legacy offender numbers to case-management person ids",
            C.F_ASSESS: "every MRAS assessment from 2017 to 2026, legacy numbers before the conversion",
            C.F_DARROW: "Darrow County Jail's booking notices of Region 4 supervisees, one row per booking charge",
            C.F_KESWICK: "Keswick County Jail's booking notices of Region 4 supervisees, one row per booking, UTC",
            C.F_HPD: "Halverton Police Department's violent-incident extract, one row per victim per offence",
            C.F_GPS: "Belvantis's event export for the Region 4 account",
            C.F_PROG: "programme referrals of Region 4 supervisees (distractor)",
            C.F_DIRECTIVE: "Department directive on GPS placement after a Category V booking",
            C.F_CODE: "Legislative Revisor's conversion table for the 2024 recodification",
            C.F_SCHED: "Belvantis contract Schedule C, billing rules",
            C.F_BRIEF: "Department research brief on statewide three-year rearrest (distractor)"}
    for f, what in made.items():
        src[f] = {"source": "synthetic (task generator)", "date": "2027-01-08", "license": "CC0-1.0", "covers": what,
                  "rows": n.get(f)}
    return {"task": "task179", "distractor_files": ["target/" + x for x in C.DISTRACTORS],
            "deliverables": [C.DELIV_XLSX, C.DELIV_DOCX], "objective": "Forecasting & Predictive Modeling",
            "domain": "Criminology & Justice / Recidivism Risk Assessment",
            "setting": ("The State of Merrow, its Department of Corrections, Region 4, Darrow and Keswick counties, the City "
                        "of Halverton and Belvantis are fictional; every record is synthetic, built by the task's generator, "
                        "and no real agency, person or case is represented. People's names are drawn from Faker's public "
                        "name lists (MIT licence)."),
            "sources": src}


def main(target, golden_dir, figures_path, submission_path, metadata_path, n=None):
    o = compute(target)
    f = figures(o)
    os.makedirs(golden_dir, exist_ok=True)
    write_workbook(os.path.join(golden_dir, C.DELIV_XLSX), o, f)
    write_memo(os.path.join(golden_dir, C.DELIV_DOCX), o, f)
    json.dump(f, open(figures_path, "w"), indent=1, sort_keys=True)
    open(submission_path, "w").write(submission(o, f))
    json.dump(metadata(n or {}), open(metadata_path, "w"), indent=2)
    return o, f
