"""Writes target/ from the clean world: the surveillance, tracing and laboratory exports, the planning note, the
situation report, the centres' capacity report, the distractors, the lookup, the dictionary and the manifest.
The setting (the Republic of Orvanthe, Dunmarrow Province, Tamberlow virus disease) is fictional."""
import csv
import datetime as dt
import io
import os
import shutil
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import world as W
import pack_data as PD
import util
import docs

HERE = os.path.dirname(os.path.abspath(__file__))
TARGET = os.path.join(HERE, "..", "target")

F_LL = "tvd_line_list_dunmarrow_2026-10-18.csv"
F_CH = "tvd_transmission_links_2026-10-18.csv"
F_CT = "contact_registry_dunmarrow_w36-w42.csv"
F_LB = "tamberlow_field_lab_pcr_w36-w42.xlsx"
F_LM = "kettlebank_phl_results_w36-w42.csv"
F_PN = "planning_note_phase3_2026-10-19.pdf"
F_SR = "phd_situation_report_w41.pdf"
F_CP = "lmr_bed_capacity_w36-w42.xlsx"
F_AD = "lmr_treatment_centre_admissions_2026-10-18.xlsx"
F_BU = "safe_burials_log_2026-10-18.csv"
F_ZS = "dunmarrow_districts.csv"
F_DD = "data_dictionary_tvd_response_2026-10.xlsx"
F_MF = "export_manifest_2026-10-19.csv"
DISTRACTORS = [F_LM, F_AD, F_BU]

LL_COLS = {"numero_cas": "case_number", "zone_sante": "district", "sexe": "sex", "age": "age", "unite_age": "age_unit",
           "date_debut_symptomes": "date_of_onset", "date_notification": "date_of_notification",
           "classification": "classification", "id_echantillon": "sample_id", "date_resultat_labo": "lab_result_date",
           "centre_prise_en_charge": "treatment_centre", "date_admission": "date_admitted", "issue": "outcome",
           "date_issue": "outcome_date", "contact_connu": "known_contact", "id_contact": "contact_id",
           "date_saisie": "date_entered"}


def csv_text(df, sep=",", crlf=True):
    bio = io.StringIO()
    df.to_csv(bio, sep=sep, index=False, lineterminator="\r\n" if crlf else "\n", quoting=csv.QUOTE_MINIMAL)
    return bio.getvalue()


def fmt_dt(x):
    return x.strftime("%d/%m/%Y %H:%M")


def style_sheet(ws, widths=None, freeze="A2"):
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    if freeze:
        ws.freeze_panes = freeze
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def build(seed=W.SEED, out=TARGET, c_off=None, l_off=None):
    P = PD.build_all(seed, c_off, l_off)
    if os.path.isdir(out):
        shutil.rmtree(out)
    os.makedirs(out)
    rng = np.random.default_rng(seed + 6000)
    # 1 line list
    ll = P["ll"]
    out_ll = ll[list(LL_COLS)].rename(columns=LL_COLS)
    out_ll["age_unit"] = out_ll["age_unit"].map({"ans": "years", "mois": "months"})
    util.write_text(os.path.join(out, F_LL), csv_text(out_ll), bom=True)
    # 2 transmission links (the tracing system's relationship export)
    ch = P["chains"]
    ch_out = pd.DataFrame({"Relationship ID": ch["relationship_id"], "Source person": ch["source_case_number"],
                           "Target person": ch["target_case_number"], "Exposure type": ch["exposure_type"],
                           "Certainty level": ch["certainty_level"], "Created on": ch["created_on"]})
    util.write_text(os.path.join(out, F_CH), csv_text(ch_out, crlf=False))
    # 3 contact registry
    rg = P["reg"]
    stat = {"Suivi termine": "Follow-up completed", "Perdu de vue": "Lost to follow-up", "Devenu cas": "Became a case",
            "En suivi": "Under follow-up", "Deplace hors zone": "Moved out of district", "Refus": "Refused"}
    rg_out = pd.DataFrame({
        "contact_id": rg["id_contact"], "follow_up_team": rg["equipe"], "district": rg["zone_sante"],
        "source_case_number": rg["numero_cas_source"], "household_code": rg["code_menage"], "initials": rg["initiales"],
        "sex": rg["sexe"], "age_years": rg["age"], "date_listed": [PD.d_fr(d) for d in rg["listing"]],
        "date_last_exposure": [PD.d_fr(d) for d in rg["last_exp"]], "follow_up_end": [PD.d_fr(d) for d in rg["end"]],
        "days_seen": rg["jours_vus"], "status": rg["statut"].map(stat), "status_date": [PD.d_fr(d) for d in rg["status_day"]],
        "case_number_if_case": rg["numero_cas_devenu"]})
    rg_out = rg_out.sort_values(["district", "contact_id"], kind="stable")
    util.write_text(os.path.join(out, F_CT), csv_text(rg_out), bom=True)
    # 4 field laboratory (analyser export, one sheet per epi week, result clock on UTC)
    lab = P["lab"]
    pulled = dt.datetime(2026, 10, 19, 8, 0)      # both laboratory exports were pulled at 08:00 on Monday 19 October
    lab = lab[lab["res_true"] <= pulled]
    wb = Workbook()
    wb.remove(wb.active)
    bk = lab[lab["lab"] == "TFL"].copy()
    bk["res_written"] = bk["res_true"] - dt.timedelta(hours=PD.UTC_OFFSET_H)
    bk["wk"] = [int((c - dt.datetime.combine(W.DAY0, dt.time())).days // 7) + 22 for c in bk["coll"]]
    for wk in range(36, 43):
        ws = wb.create_sheet("W%d" % wk)
        ws.append(["Sample no.", "Case / alert no.", "Reason", "Collected", "Received at lab", "Result time (analyser)",
                   "Result", "Ct"])
        sub = bk[bk["wk"] == wk].sort_values(["rec", "id_echantillon"], kind="stable")
        for r in sub.itertuples():
            ct = "" if r.resultat != "Positive" else round(float(rng.uniform(17.5, 34.0)), 1)
            ws.append([r.id_echantillon, r.numero, r.motif, fmt_dt(r.coll), fmt_dt(r.rec), fmt_dt(r.res_written),
                       r.resultat, ct])
        style_sheet(ws, [16, 18, 18, 17, 17, 21, 13, 7])
    util.save_workbook(wb, os.path.join(out, F_LB), title="Tamberlow field laboratory - PCR analyser export")
    # 5 public health laboratory LIMS export (distractor)
    mb = lab[lab["lab"] == "KPL"].sort_values(["rec", "id_echantillon"], kind="stable")
    mb_out = pd.DataFrame({"sample_code": mb["id_echantillon"], "patient": mb["numero"], "reason": mb["motif"],
                           "collected_at": [x.strftime("%Y-%m-%d %H:%M") for x in mb["coll"]],
                           "received_at": [x.strftime("%Y-%m-%d %H:%M") for x in mb["rec"]],
                           "validated_at": [x.strftime("%Y-%m-%d %H:%M") for x in mb["res_true"]],
                           "result": mb["resultat"]})
    util.write_text(os.path.join(out, F_LM), csv_text(mb_out, sep=";"))
    docs.write_all(P, out, rng)
    return P


if __name__ == "__main__":
    build()
    for f in sorted(os.listdir(TARGET)):
        print(f, os.path.getsize(os.path.join(TARGET, f)))
