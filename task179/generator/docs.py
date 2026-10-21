"""Documents and small tables of the pack (fictional setting: Republic of Orvanthe, Dunmarrow Province)."""
import datetime as dt
import os
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, Spacer, Table, TableStyle
import world as W
import util

BASE = ParagraphStyle("b", fontName="Helvetica", fontSize=9.6, leading=13.2, spaceAfter=6)
SMALL = ParagraphStyle("s", parent=BASE, fontSize=7.8, leading=10)
H1 = ParagraphStyle("h1", parent=BASE, fontName="Helvetica-Bold", fontSize=13, leading=16, spaceAfter=8)
H2 = ParagraphStyle("h2", parent=BASE, fontName="Helvetica-Bold", fontSize=10.4, leading=13.5, spaceBefore=6, spaceAfter=4)
HEAD = ParagraphStyle("hd", parent=BASE, fontSize=8.4, leading=11, textColor=colors.HexColor("#333333"))

DISTRICT_POP = {"Tamberlow": 172000, "Quillmarsh": 141000, "Threlmarren": 148000, "Wendlecombe": 96000,
                "Brindlecourt": 118000, "Kettlebank North": 412000, "Kettlebank South": 398000}
DISTRICT_CODE = {"Tamberlow": "DMW-TAM", "Quillmarsh": "DMW-QUI", "Threlmarren": "DMW-VOS", "Wendlecombe": "DMW-ORD",
                 "Brindlecourt": "DMW-BRI", "Kettlebank North": "DMW-KBN", "Kettlebank South": "DMW-KBS"}
DISTRICT_AREAS = {"Tamberlow": 17, "Quillmarsh": 15, "Threlmarren": 14, "Wendlecombe": 12, "Brindlecourt": 10,
                  "Kettlebank North": 18, "Kettlebank South": 17}


def fix_pdf(data, stamp="20261019081200"):
    data = data.replace(b"20000101000000", stamp.encode())
    sig = b" ReportLab Generated PDF document (opensource)"
    rep = b" Larrowmede Medical Relief, epidemiology cell"
    data = data.replace(sig, rep + b" " * (len(sig) - len(rep)))
    sig2 = b"% ReportLab generated PDF document -- digest (opensource)"
    rep2 = b"% LMR epidemiology cell"
    data = data.replace(sig2, rep2 + b" " * (len(sig2) - len(rep2)))
    return data.replace(b"+00'00'", b"+01'00'")


def footer(text):
    def fn(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.4)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(52, 30, text)
        canvas.drawRightString(595 - 52, 30, "%d" % doc.page)
        canvas.restoreState()
    return fn


def tbl(rows, widths, size=8.2):
    t = Table(rows, colWidths=widths, repeatRows=1)
    st = [("FONT", (0, 0), (-1, -1), "Helvetica", size), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", size),
          ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.black),
          ("LINEBELOW", (0, -1), (-1, -1), 0.4, colors.grey), ("TOPPADDING", (0, 0), (-1, -1), 2),
          ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]
    for c in range(1, len(rows[0])):
        st.append(("ALIGN", (c, 1), (c, -1), "RIGHT"))
    t.setStyle(TableStyle(st))
    return t


def planning_note(P, path):
    s = []
    s.append(Paragraph("Larrowmede Medical Relief &nbsp;|&nbsp; Dunmarrow TVD response &nbsp;|&nbsp; Epidemiology cell", HEAD))
    s.append(Paragraph("Planning note PN-DMW-07 &nbsp;&nbsp; 19 October 2026 &nbsp;&nbsp; Owner: Jacqueline Arnold (response "
                       "epidemiologist) &nbsp;&nbsp; To: emergency desk; copy Tanner Molina, Kristin Sampson", HEAD))
    s.append(Spacer(1, 6))
    s.append(Paragraph("December's isolation capacity at Kettlebank North and Tamberlow is set on the cases of weeks 49 to 52", H1))
    s.append(Paragraph("1. What the desk commits", H2))
    s.append(Paragraph(
        "The emergency desk orders beds, tents, staff rotations and consumables for December in one round, closing on "
        "30 October, and it orders against a single planning figure: the number of new cases of Tamberlow virus disease, "
        "confirmed and probable, with symptom onset in epidemiological weeks 49 to 52 (Monday 30 November to Sunday 27 "
        "December 2026) in the seven districts of Dunmarrow Province with cases (Tamberlow, Quillmarsh, Threlmarren, Wendlecombe, "
        "Brindlecourt, Kettlebank North and Kettlebank South). The centres' share of those cases, the admission "
        "rate and the length of stay are applied by the desk afterwards; they are not part of the planning figure.", BASE))
    s.append(Paragraph("2. The scenario", H2))
    s.append(Paragraph(
        "Phase 3 of the provincial response plan starts on Monday 2 November: contact tracing extended to contacts of "
        "contacts in every district, the directorate's third treatment centre opened at Threlmarren, and the safe burial "
        "teams doubled. The coordination meeting of 16 October adopted three planning scenarios for phase 3 and asked "
        "partners to plan on the central one, which is the scenario this figure is set on: every case with symptom onset "
        "on or after 2 November causes half as many new cases as cases did over weeks 36 to 41 (31 August to 11 October), "
        "and the reproduction number of those six weeks is the one implied by the exponential growth rate of their daily "
        "onsets. The lower and upper "
        "scenarios (a reduction of 35 and of 65 per cent) are carried in the provincial plan and are not used here.", BASE))
    s.append(Paragraph("3. Records", H2))
    s.append(Paragraph(
        "The figure rests on this outbreak's own records: the provincial line list and the response's transmission links "
        "as extracted on Sunday 18 October at 23:59. The provincial protocol closes notification within four days of "
        "onset, so onsets to Sunday 11 October are complete in the extract. Tamberlow virus disease has not been "
        "described outside this outbreak, so no parameter comes from elsewhere. Cases are counted under the provincial "
        "case definition (confirmed or probable as the line list classifies them); people found not to be cases and "
        "suspected cases awaiting a result are not cases.", BASE))
    s.append(Paragraph("4. What goes to the desk beside the figure", H2))
    s.append(Paragraph(
        "Two checks on the response the scenario leans on go with the figure: the follow-up of contacts in the four "
        "districts where our tracing teams work (Tamberlow, Quillmarsh, Kettlebank North and Brindlecourt), and the "
        "turnaround of the Tamberlow field laboratory, which tests every sample from our Tamberlow centre and from the "
        "Tamberlow and Quillmarsh alert teams; the response's target is a final result no more than 24 hours after "
        "collection. A person counts as having completed follow-up when the 21-day follow-up "
        "that ends last for them is recorded as completed by either team.", BASE))
    rows = [["Phase 3 component", "Lead", "From", "Status at 19 Oct"],
            ["Tracing of contacts of contacts, all districts", "PHD / partners", "2 Nov", "Teams recruited"],
            ["Threlmarren treatment centre (directorate), 60 beds", "PHD", "2 Nov", "Construction"],
            ["Safe burial teams, 14 to 28", "PHD / volunteers", "26 Oct", "Recruiting"],
            ["Community isolation, Quillmarsh and Wendlecombe", "LMR / PHD", "9 Nov", "Sites identified"]]
    s.append(tbl(rows, [210, 90, 50, 110]))
    s.append(Paragraph("Source: provincial response plan, phase 3 (coordination meeting, 16 October 2026).", SMALL))
    s.append(Spacer(1, 6))
    s.append(Paragraph("The figure is reviewed when the extract of 1 November is in. Weeks are ISO weeks, Monday to Sunday, "
                       "as in the directorate's reports.", SMALL))

    def build(doc):
        doc.build(s, onFirstPage=footer("LMR Dunmarrow - PN-DMW-07 - internal"),
                  onLaterPages=footer("LMR Dunmarrow - PN-DMW-07 - internal"))
    util.write_bytes(path, fix_pdf(util.pdf_bytes(build, "PN-DMW-07 Planning figure weeks 49-52")))


def ndays(x):
    return (dt.datetime.strptime(x, "%d/%m/%Y").date() - W.DAY0).days


def sitrep(P, path):
    ll = P["ll"]
    cases = ll[ll["_kind"].isin(["case", "dup"])].copy()
    cases["nday"] = [ndays(x) for x in cases["date_notification"]]
    cases["nwk"] = cases["nday"] // 7 + 22
    cases = cases[cases["nday"] <= 141]
    s = [Paragraph("Dunmarrow Provincial Health Directorate &nbsp;|&nbsp; Tamberlow virus disease", HEAD),
         Paragraph("Situation report 14, epidemiological week 41 (5 to 11 October 2026)", H1),
         Paragraph("New cases are counted by week of notification, confirmed and probable, as notified when the report "
                   "closed (Tuesday 13 October 2026). Figures for the latest week are provisional.", BASE)]
    rows = [["District"] + ["W%d" % w for w in range(36, 42)] + ["Total"]]
    tot = [0] * 7
    for z in W.ZONES:
        sub = cases[cases["zone_sante"] == z]
        r = [z]
        for i, w in enumerate(range(36, 42)):
            v = int((sub["nwk"] == w).sum())
            r.append(v)
            tot[i] += v
        r.append(int(len(sub)))
        tot[6] += int(len(sub))
        rows.append(r)
    rows.append(["Province"] + tot)
    s.append(tbl(rows, [92] + [42] * 6 + [52]))
    s.append(Paragraph("Source: provincial line list, extract of 13 October 2026.", SMALL))
    deaths = int((cases["issue"] == "Died").sum())
    s.append(Paragraph(
        "Deaths among the cases notified by 13 October: %d. Four treatment centres are open (Kettlebank North and Tamberlow, "
        "Larrowmede Medical Relief; Kettlebank General isolation unit and Threlmarren, the directorate); the directorate's "
        "third centre opens at Threlmarren on 2 November. Both laboratories (the Kettlebank public health laboratory and the "
        "Tamberlow field laboratory) have worked around the clock since week 38." % deaths, BASE))
    s.append(Paragraph("Points to note", H2))
    s.append(Paragraph(
        "Larrowmede Medical Relief's tracing teams took over follow-up in Tamberlow and Quillmarsh from week 37 and support "
        "the directorate's teams in Kettlebank North and Brindlecourt. Phase 3 extends tracing to contacts of contacts.", BASE))

    def build(doc):
        doc.build(s, onFirstPage=footer("Dunmarrow PHD - TVD situation report 14"),
                  onLaterPages=footer("Dunmarrow PHD - TVD situation report 14"))
    util.write_bytes(path, fix_pdf(util.pdf_bytes(build, "TVD situation report 14 W41", author="Dunmarrow PHD"), "20261013180500"))


def capacity(P, path, rng):
    wb = Workbook()
    ws = wb.active
    ws.title = "Beds by week"
    ws.append(["Week", "Monday", "Centre", "Beds open (suspected)", "Beds open (confirmed)", "Beds staffed",
               "Nurses on roster", "Doctors on roster"])
    base = {"Kettlebank North TC": (24, 60), "Tamberlow TC": (30, 90)}
    for wk in range(36, 43):
        mon = W.date_of((wk - 22) * 7)
        for c, (sb, k) in base.items():
            ks = k + (wk - 36) * (6 if c == "Tamberlow TC" else 4)
            staffed = ks + sb - int(rng.integers(4, 15))
            ws.append([wk, mon.strftime("%d/%m/%Y"), c, sb + (wk - 36) * 2, ks, staffed,
                       int(staffed / 2.6) + int(rng.integers(0, 4)), int(staffed / 14) + 1])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True)
    ws.freeze_panes = "A2"
    for col, w in zip("ABCDEFGH", [8, 11, 20, 12, 12, 12, 11, 11]):
        ws.column_dimensions[col].width = w
    util.save_workbook(wb, path, title="LMR bed capacity W36-W42")


def admissions(P, path, rng):
    cases = P["cases"]
    wb = Workbook()
    wb.remove(wb.active)
    for c, code in (("Kettlebank North TC", "KN"), ("Tamberlow TC", "TB")):
        ws = wb.create_sheet(c)
        ws.append(["Patient no.", "Case no.", "Admitted", "Ward", "Lab result", "Discharged", "Discharge reason"])
        sub = cases[(cases["centre"] == c) & (cases["adm_day"] >= 0) & (cases["adm_day"] <= W.EXTRACT_DAY)].sort_values("adm_day", kind="stable")
        k = 0
        for r in sub.itertuples():
            k += 1
            a = dt.datetime.combine(W.date_of(r.adm_day), dt.time()) + dt.timedelta(hours=float(rng.uniform(0, 24)))
            out = r.out_day if r.out_day <= W.EXTRACT_DAY else None
            dis = (dt.datetime.combine(W.date_of(out), dt.time()) + dt.timedelta(hours=float(rng.uniform(6, 20)))).strftime("%d/%m/%Y %H:%M") if out is not None else ""
            ws.append(["%s-%05d" % (code, 100 + 2 * k), r.numero_cas, a.strftime("%d/%m/%Y %H:%M"), "Confirmed", "Positive", dis,
                       ("Died" if r.died else "Recovered") if out is not None else ""])
        for c2 in ws[1]:
            c2.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for col, w in zip("ABCDEFG", [11, 17, 16, 11, 11, 16, 15]):
            ws.column_dimensions[col].width = w
    util.save_workbook(wb, path, title="LMR treatment centre admissions")


def burials(P, path, rng):
    import pack
    cases = P["cases"]
    rows = []
    k = 0
    for r in cases[(cases["died"]) & (cases["out_day"] <= W.EXTRACT_DAY) & (cases["out_day"] >= 49)].sort_values("out_day", kind="stable").itertuples():
        if not r.probable and r.centre and rng.random() < 0.7:
            continue
        k += 1
        d = int(r.out_day + rng.integers(0, 3))
        if d > W.EXTRACT_DAY:
            continue
        rows.append({"burial_id": "SB-%05d" % (300 + k), "case_number": r.numero_cas if rng.random() < 0.93 else "",
                     "district": r.zone, "date_alert": W.date_of(r.out_day).strftime("%d/%m/%Y"),
                     "date_burial": W.date_of(d).strftime("%d/%m/%Y"), "swab_taken": "Yes" if rng.random() < 0.8 else "No",
                     "team": "SB-%02d" % int(rng.integers(1, 15)), "family_consent": "Yes" if rng.random() < 0.96 else "No"})
    util.write_text(path, pack.csv_text(pd.DataFrame(rows)), bom=True)


def districts(path):
    import pack
    rows = [{"district_code": DISTRICT_CODE[z], "district": z, "health_areas": DISTRICT_AREAS[z],
             "population_2026": DISTRICT_POP[z]} for z in W.ZONES]
    util.write_text(path, pack.csv_text(pd.DataFrame(rows)), bom=True)


def dictionary(path, files):
    wb = Workbook()
    ws = wb.active
    ws.title = "Files"
    ws.append(["File", "System", "One row is", "Notes"])
    for f in files:
        ws.append(list(f))
    defs = {
        "Line list": [
            ("case_number", "Case number given at notification: province, district code, year and serial in the notifying district."),
            ("classification", "Classification at the extract under the provincial case definition: Confirmed, Probable, Not a case (a notified person later shown not to be a case) or Suspected (result awaited)."),
            ("date_of_onset", "Date of symptom onset reported at the case investigation. Rows entered from the field laboratory's forms carry ISO dates."),
            ("date_of_notification", "Date the case was notified to the district surveillance team."),
            ("sample_id", "Laboratory number of the confirming sample (KPL: Kettlebank public health laboratory, TFL: Tamberlow field laboratory)."),
            ("lab_result_date", "Date the laboratory telephoned the result to the district."),
            ("known_contact / contact_id", "Whether the case was on a contact list before onset, and the listing's id in the contact registry."),
            ("outcome / outcome_date", "Outcome known at the extract."),
        ],
        "Transmission links": [
            ("Source person / Target person", "Case numbers of the source case and of the case it is linked to. A link is entered by the tracing team when the target case is notified and its investigation names the source."),
            ("Exposure type", "Exposure the investigation recorded for the link."),
            ("Certainty level", "Tracing team's grading of the link: High, Medium, Low."),
            ("Created on", "Date the link was entered."),
        ],
        "Contact registry": [
            ("contact_id", "Listing id; prefix LMR (tablet teams of Larrowmede Medical Relief) or PHD (directorate teams, paper forms entered at the district office)."),
            ("source_case_number", "Case whose investigation listed the contact. A person exposed to two cases is listed once for each."),
            ("household_code / initials / sex / age_years", "Household of the contact, initials, sex and age as recorded at listing."),
            ("date_last_exposure / follow_up_end", "Last exposure to the source case, and the end of the 21-day follow-up from it."),
            ("days_seen", "Days on which the contact was seen by the team that holds the listing."),
            ("status", "Status at the extract: Follow-up completed, Under follow-up, Lost to follow-up, Moved out of district, Refused, Became a case."),
        ],
        "Tamberlow field laboratory": [
            ("Collected / Received at lab", "Collection time from the sample form, reception time keyed at the laboratory bench."),
            ("Result time (analyser)", "Time the PCR analyser wrote the result. The TVD cartridge runs 92 minutes from loading."),
            ("Reason", "Diagnostic, Discharge control (sample taken before a treated patient is discharged)."),
            ("Result", "Positive, Negative, Indeterminate (the sample is re-run and its final result follows on a new row)."),
        ],
    }
    for name, rows in defs.items():
        sh = wb.create_sheet(name)
        sh.append(["Field", "Definition"])
        for r in rows:
            sh.append(list(r))
        sh.column_dimensions["A"].width = 36
        sh.column_dimensions["B"].width = 110
        for c in sh[1]:
            c.font = Font(bold=True)
    for col, w in zip("ABCD", [44, 40, 36, 52]):
        ws.column_dimensions[col].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
    util.save_workbook(wb, path, title="Data dictionary TVD response Dunmarrow")


def write_all(P, out, rng):
    import pack as PK
    planning_note(P, os.path.join(out, PK.F_PN))
    sitrep(P, os.path.join(out, PK.F_SR))
    capacity(P, os.path.join(out, PK.F_CP), rng)
    admissions(P, os.path.join(out, PK.F_AD), rng)
    burials(P, os.path.join(out, PK.F_BU), rng)
    districts(os.path.join(out, PK.F_ZS))
    files = [
        (PK.F_LL, "Provincial surveillance database", "a notified person", "Extract of 18/10/2026 23:59"),
        (PK.F_CH, "Outbreak tracing system, relationship export", "a source-to-case link", "Extract of 18/10/2026 23:59"),
        (PK.F_CT, "Contact registry (LMR tablets, PHD paper forms)", "a listing of a person as a contact", "Listings from 31/08/2026"),
        (PK.F_LB, "PCR analyser export, Tamberlow field laboratory", "a test result", "One sheet per epidemiological week; pulled 19/10/2026 08:00"),
        (PK.F_LM, "LIMS, Kettlebank public health laboratory", "a validated result", "Weeks 36 to 42; pulled 19/10/2026 08:00"),
        (PK.F_PN, "LMR epidemiology cell", "", "Planning note PN-DMW-07"),
        (PK.F_SR, "Dunmarrow Provincial Health Directorate", "", "Situation report 14"),
        (PK.F_CP, "LMR logistics", "a centre in a week", "Beds and rosters"),
        (PK.F_AD, "Patient register, LMR treatment centres", "an admission", "Confirmed ward"),
        (PK.F_BU, "Safe burial teams", "a burial", "From 13/07/2026"),
        (PK.F_ZS, "Dunmarrow Provincial Health Directorate", "a district", "Population projections 2026"),
    ]
    dictionary(os.path.join(out, PK.F_DD), files)
    man = pd.DataFrame([{"file": f, "system": sy, "one_row_is": l, "note": n} for f, sy, l, n in files] +
                       [{"file": PK.F_DD, "system": "LMR epidemiology cell", "one_row_is": "", "note": "Field definitions"}])
    man.insert(1, "compiled_on", "19/10/2026 08:12")
    man.insert(2, "compiled_by", "LMR epidemiology cell (J. Arnold)")
    util.write_text(os.path.join(out, PK.F_MF), PK.csv_text(man), bom=True)
