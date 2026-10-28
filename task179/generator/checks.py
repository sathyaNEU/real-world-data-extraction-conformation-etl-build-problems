"""task179 generator assertions, run on the cut pack under target/ with the generator's own reference solution.

Blocks: the call and its family (the ladder, the standing family, the bin), the three rules (R1 to R3c), the semantic
fork table (SF1 to SF13), mess neutrality per Layer 1 item, the asks (the cracker sheet and the separation line), the
realism sweep, the input gates and the distractors.

    python3 checks.py <target>"""
from __future__ import annotations

import datetime as dt
import json
import os
import re
import sys
import zipfile
from collections import Counter, defaultdict

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import calc          # noqa: E402
import codes as K    # noqa: E402
import common as C   # noqa: E402
import docs as DC    # noqa: E402

RES = []
D = dt.date
LO, HI = 450.0, 550.0          # the golden's printed cell at the nearest 100
Q = C.ROUND_CALL / 4.0         # a quarter of the rounding unit: every route sits this far from a cell edge


def check(name, cond, detail=""):
    RES.append((name, bool(cond), detail))
    print("  %-4s %-66s %s" % ("PASS" if cond else "FAIL", name, detail))


def printed(v):
    return C.nearest(v, C.ROUND_CALL)


def pdf_text(path):
    from pypdf import PdfReader
    return " ".join(" ".join((p.extract_text() or "").split()) for p in PdfReader(path).pages)


def main(target):
    figs = json.load(open(os.path.join(HERE, "figures.json")))
    data = calc.load(target)
    sol = calc.Solution(target, data=data)
    forms = sol.all_forms()
    golden = forms["victim 12 months"][0]
    default = forms["own attributes"][0]
    f26, tr = sol.f26, sol.train

    print("the call")
    check("call: the golden form prints the call", printed(golden) == figs["call_forecast_2027"] == 500, "%.1f" % golden)
    check("call: the golden sits a quarter unit clear of its cell's edges", C.edge_distance(golden, C.ROUND_CALL) >= Q,
          "edge distance %.1f" % C.edge_distance(golden, C.ROUND_CALL))
    linked = {k: v[0] for k, v in forms.items() if k != "own attributes"}
    check("standing family: every linked form sits a quarter unit inside the call's cell", all(LO + Q <= v <= HI - Q for v in linked.values()),
          ", ".join("%s %.1f" % (k, v) for k, v in linked.items()))
    spread = max(linked.values()) - min(linked.values())
    check("standing family: spread under a third of the cell", spread < C.ROUND_CALL / 3.0, "%.1f" % spread)
    aic = {k: v[1] for k, v in forms.items()}
    best = min(aic, key=aic.get)
    check("the AIC-best form is a linked form and prints the call", best != "own attributes" and LO <= forms[best][0] < HI,
          "%s %.1f" % (best, forms[best][0]))
    check("the own-attribute form loses to every linked form by more than 100 AIC",
          all(aic["own attributes"] - aic[k] > 100 for k in linked), "%.1f" % min(aic["own attributes"] - aic[k] for k in linked))

    # ------------------------------------------------------------ the ladder: cheap routes off the call
    print("the ladder")
    rungs = {}
    rungs["0 own attributes"] = default
    tr24 = pd.concat([sol.frames[Y] for Y in calc.YEARS if Y < 2025], ignore_index=True)
    v25, _ = sol.forecast("own attributes", train=tr24, f=sol.frames[2025])
    obs25 = float(sol.frames[2025]["y"].sum())
    rungs["1 recalibrated on the latest back-test"] = default * obs25 / v25
    rungs["2 latest year's rate on the 2026 caseload"] = obs25 / len(sol.frames[2025]) * len(f26)
    ys = np.array([sol.frames[Y]["y"].sum() / len(sol.frames[Y]) for Y in calc.YEARS])
    tt = np.array(calc.YEARS, float)
    b1, b0 = np.polyfit(tt, ys, 1)
    rungs["3 linear trend in the yearly rate"] = (b0 + b1 * 2026) * len(f26)
    h = data["hpd"]
    hi_ = h[(h["victim_type"] == "I") & (h["supplement_no"] == "0")].drop_duplicates(["incident_number", "victim_seq"])
    yr = hi_["occurred_date"].str[6:10].astype(int).value_counts()
    rungs["4 scaled by the city's 2026 victims"] = default * yr[2026] / yr[2025]
    trc = tr.copy()
    trc["cityv"] = trc["Y"].map(lambda Y: yr.get(Y, np.nan) / 1000.0) * trc["city"]
    f26c = f26.copy()
    f26c["cityv"] = yr[2026] / 1000.0 * f26c["city"]
    m = calc.fit(trc, calc.BASE + ["cityv"])
    rungs["5 city violence as a covariate"] = calc.predict(m, f26c, calc.BASE + ["cityv"])
    from sklearn.ensemble import HistGradientBoostingClassifier
    gb = HistGradientBoostingClassifier(max_iter=200, learning_rate=0.05, max_leaf_nodes=15, random_state=0)
    gb.fit(tr[calc.BASE].astype(float), tr["y"].astype(int))
    rungs["6 boosted trees on own attributes"] = float(gb.predict_proba(f26[calc.BASE].astype(float))[:, 1].sum())
    for k, v in rungs.items():
        check("rung %s prints 400, a quarter unit below the call's cell" % k, printed(v) == 400 and v <= LO - Q,
              "%.1f prints %d" % (v, printed(v)))
    near = max(rungs.values())
    check("nearest wrong rung clears the golden cell by a quarter unit", LO - near >= Q, "%.1f below %.0f" % (LO - near, LO))
    gb2 = HistGradientBoostingClassifier(max_iter=200, learning_rate=0.05, max_leaf_nodes=15, random_state=0)
    gb2.fit(tr[calc.BASE + ["v12"]].astype(float), tr["y"].astype(int))
    gbl = float(gb2.predict_proba(f26[calc.BASE + ["v12"]].astype(float))[:, 1].sum())
    check("boosted trees with the victim record print the call", LO + Q <= gbl <= HI - Q, "%.1f" % gbl)

    # ------------------------------------------------------------ the margin over the judge's estimators and partial routes
    print("margins")
    def flags(f):
        f = f.copy()
        son, rec = [], []
        for p, Y in zip(f["pid"], f["Y"]):
            ds = sol.links.get(p, [])
            son.append(int(any(D(Y, 9, 1) <= x <= D(Y, 11, 30) for x in ds)))
            last = max([x for x in ds if x <= D(Y, 12, 31)], default=None)
            rec.append(-1 if last is None else (D(Y, 12, 31) - last).days)
        f["son"], f["rec"] = son, rec
        return f
    trs, f26s = flags(tr), flags(f26)
    rv_, rn_ = tr.loc[tr["v12"] == 1, "y"].mean(), tr.loc[tr["v12"] == 0, "y"].mean()
    est = {"victim-stratified rates": rv_ * f26["v12"].sum() + rn_ * (len(f26) - f26["v12"].sum()),
           "a logistic model with a September to November victim flag":
               calc.predict(calc.fit(trs, calc.BASE + ["v12", "son"]), f26s, calc.BASE + ["v12", "son"])}
    def bucket(v):
        return "none" if v < 0 else "0-91" if v <= 91 else "92-182" if v <= 182 else "183-365" if v <= 365 else \
            "366-730" if v <= 730 else "731+"
    trs["b"], f26s["b"] = trs["rec"].map(bucket), f26s["rec"].map(bucket)
    rates = trs.groupby("b")["y"].mean()
    est["rates stratified by the time since the last victim record"] = float(sum(rates.get(b, rn_) for b in f26s["b"]))
    tr4 = pd.concat([sol.frames[Y] for Y in (2022, 2023, 2024, 2025)], ignore_index=True)
    est["the linked model on the 2022 to 2025 year ends only"] = calc.predict(calc.fit(tr4, calc.BASE + ["v12"]), f26,
                                                                              calc.BASE + ["v12"])
    est["a six-month victim flag only"] = calc.predict(calc.fit(tr, calc.BASE + ["v6"]), f26, calc.BASE + ["v6"])
    for k, v in est.items():
        check("margin: %s sits a quarter unit inside the call's cell" % k, LO + Q <= v <= HI - Q, "%.1f" % v)
    v24 = calc.predict(calc.fit(tr, calc.BASE + ["v24"]), f26, calc.BASE + ["v24"])
    check("margin: a flag that keeps only the 12 to 24 month records prints 400, a quarter unit below the cell",
          printed(v24) == 400 and v24 <= LO - Q, "%.1f" % v24)

    # ------------------------------------------------------------ back-tests
    print("back-tests")
    for Yt in (2023, 2024, 2025):
        trb = pd.concat([sol.frames[Y] for Y in calc.YEARS if Y < Yt], ignore_index=True)
        vo, _ = sol.forecast("own attributes", train=trb, f=sol.frames[Yt])
        vl, _ = sol.forecast("victim 12 months", train=trb, f=sol.frames[Yt])
        ob = float(sol.frames[Yt]["y"].sum())
        check("back-test %d: own-attribute model within 15 per cent" % Yt, abs(vo / ob - 1) < 0.15, "%.1f against %d" % (vo, ob))
        check("back-test %d: linked model within 15 per cent" % Yt, abs(vl / ob - 1) < 0.15, "%.1f against %d" % (vl, ob))
    trb = pd.concat([sol.frames[Y] for Y in calc.YEARS if Y < 2025], ignore_index=True)
    vl25, _ = sol.forecast("victim 12 months", train=trb, f=sol.frames[2025])
    check("the golden recalibrated on the latest back-test still prints the call", LO <= golden * obs25 / vl25 < HI,
          "%.1f" % (golden * obs25 / vl25))

    # ------------------------------------------------------------ the record's own control
    print("the record's own control")
    vr = tr.loc[tr["v12"] == 1, "y"].mean()
    nr = tr.loc[tr["v12"] == 0, "y"].mean()
    check("R1 victims' booking rate at least three times the rest's", vr > 3 * nr, "%.3f against %.3f" % (vr, nr))
    m0 = calc.fit(tr, calc.BASE)
    import statsmodels.api as sm
    p0 = m0.predict(sm.add_constant(tr[calc.BASE].astype(float), has_constant="add"))
    under = []
    for Y in calc.YEARS:
        sel = (tr["Y"] == Y) & (tr["v12"] == 1)
        under.append(tr.loc[sel, "y"].sum() / p0[sel].sum())
    check("the own-attribute model under-predicts the victims in every training year", min(under) > 1.5,
          ", ".join("%.2f" % u for u in under))
    autumn = sum(1 for p in f26.loc[f26["v12"] == 1, "pid"]
                 if any(x.year == 2026 and 9 <= x.month <= 11 for x in sol.links.get(p, [])))
    check("memo: most of the 2026 victims on the caseload were recorded from September to November",
          autumn > 0.5 * figs["cc_victims_on_caseload_2026"], "%d of %d" % (autumn, figs["cc_victims_on_caseload_2026"]))
    check("2026 victims share at least twice 2025's", f26["v12"].mean() > 2 * sol.frames[2025]["v12"].mean(),
          "%.3f against %.3f" % (f26["v12"].mean(), sol.frames[2025]["v12"].mean()))
    check("figures: the caseload, the victims, the rates", (len(f26) == figs["cc_caseload_2026"] and
          int(f26["v12"].sum()) == figs["cc_victims_on_caseload_2026"] and
          int(sol.frames[2025]["v12"].sum()) == figs["cc_victims_on_caseload_2025"] and
          abs(100 * vr - figs["cc_rate_victims_pct"]) < 0.051 and abs(100 * nr - figs["cc_rate_others_pct"]) < 0.051))

    # ------------------------------------------------------------ R1 to R3
    print("the three rules")
    cols = [c for f in (C.F_EPIS, C.F_DARROW, C.F_ASSESS, C.F_XWALK) for c in data[{C.F_EPIS: "epi", C.F_DARROW: "dj",
                                                                                    C.F_ASSESS: "ras", C.F_XWALK: "xw"}[f]].columns]
    check("R1 no supervision, jail or assessment column carries victimization", not any(re.search(r"victim", c, re.I) for c in cols),
          "%d columns" % len(cols))
    check("R1 the police file carries no state number or case-management id", not any(c.lower() in ("sid", "cms_person_id")
                                                                                        for c in data["hpd"].columns))
    check("R1 the decisive computation reads the police rows", len(sol.links) > 1000, "%d people linked" % len(sol.links))
    check("R3a the default path files the named wrong answer (400)", printed(default) == 400, "%.1f" % default)
    check("R3b the default path with only the victim record added files the golden", printed(golden) == 500)
    sup_rows = h[(h["victim_type"] == "I")]
    share = len(sol.links) / max(1, sup_rows["victim_name"].nunique())
    check("R3c the police rows the step uses are a small minority of the file", share < 0.10, "%.3f" % share)
    mm = hi_.assign(ym=hi_["occurred_date"].str[6:10] + hi_["occurred_date"].str[0:2]).groupby("ym").size()
    check("R3c no repeated monthly count block in the police file", mm.duplicated().sum() < 12, "%d repeats" % mm.duplicated().sum())

    # ------------------------------------------------------------ semantic forks and mess, each through the solution
    print("semantic fork table and mess neutrality")

    def alt(**kw):
        s2 = calc.Solution(target, data=data, **kw)
        return s2.forecast("victim 12 months")[0], s2.forecast("own attributes")[0]

    variants = {
        "SF1 an episode open until (not through) its close date": dict(inclusive=False),
        "SF3 Keswick's clock left in UTC": dict(utc=False),
        "SF4 every booking type read as a new charge": dict(new_only=False),
        "SF6 bookings after supervision ended kept": dict(under_only=False),
        "SF9 victimization dated by the report": dict(victim_date="report"),
        "SF10 names matched with nicknames": dict(fuzzy=True),
        "SF11 the 2020 outcome year left out": dict(years=[Y for Y in calc.YEARS if Y != 2019]),
        "SF12 MRAS totals standardised within version": dict(std_score=True),
        "M1 Darrow charge rows read as bookings": dict(group_rows=False),
        "M2 the crosswalk's merged duplicates left unmapped": dict(xwalk="first"),
        "SF2 people counted per open episode": dict(persons_once=False),
        "M3 a county transfer read as a new spell": dict(merge_transfers=False),
        "M4 names matched as written, without cleaning": dict(clean_names=False),
    }
    for k, kw in variants.items():
        g, d0 = alt(**kw)
        check("%s: same call, a quarter unit clear" % k, LO + Q <= g <= HI - Q and printed(d0) == 400 and d0 <= LO - Q,
              "golden %.1f, default %.1f" % (g, d0))
    # the critical component's victim counts: step 3 files the match key and the date field, so the literal count is
    # the golden's; stated to the nearest 10, each sits clear of its rounding edge
    a26, a25 = int(f26["v12"].sum()), int(sol.frames[2025]["v12"].sum())
    check("component: the caseload's victims on the filed key, to the nearest 10, sit clear of the rounding edge",
          C.nearest(a26, 10) == figs["cc_victims_on_caseload_2026_about"] and
          C.nearest(a25, 10) == figs["cc_victims_on_caseload_2025_about"] and
          min(abs(a26 % 10 - 5), abs(a25 % 10 - 5)) >= 2, "%d and %d" % (a26, a25))
    # SF5 former sections: the table against the section titles
    check("SF5 every former section in the bookings converts through the Revisor's table",
          all(c in K.OLD_TO_NEW or c in K.NEW or c == "" for c in set(data["dj"]["statute"])))
    # the Revisor's table covers the former titles, not only the sections the records cite
    cited = set(data["dj"]["statute"]) | set(data["epi"]["lead_offense_statute"]) | set(
        p_.split(" ")[0] for ch in data["kj"]["Charges"] for p_ in str(ch).split(" | "))
    extra_old, extra_new = set(r[0] for r in K.EXTRA_ROWS), set(r[1] for r in K.EXTRA_ROWS)
    check("the conversion table converts sections no record cites, none of them to Category V",
          len(K.EXTRA_ROWS) >= 40 and not (extra_old | extra_new) & cited and not extra_new & set(K.SCHEDULE_B),
          "%d of %d rows uncited" % (len(K.EXTRA_ROWS), len(K.EXTRA_ROWS) + len(K.ROWS)))
    # SF7 both jails
    check("SF7 both jails carry Category V new arrests of supervisees", sol.panel.nb.groupby("jail")["schedb"].sum().min() > 100)
    # SF8 the count: expected persons against a rounded draw of each person
    draws = []
    m1 = calc.fit(tr, calc.BASE + ["v12"])
    pr = m1.predict(sm.add_constant(f26[calc.BASE + ["v12"]].astype(float), has_constant="add")).values
    rng = np.random.default_rng(7)
    for _ in range(200):
        draws.append((rng.random(len(pr)) < pr).sum())
    check("SF8 the median of simulated counts prints the call", LO <= np.median(draws) < HI, "%.0f" % np.median(draws))
    check("SF13 the population is the year-end caseload, not 2027's intakes", len(f26) == figs["cc_caseload_2026"])

    # ------------------------------------------------------------ asks
    print("asks")
    import ceiling as CE
    out = CE.main(target)
    check("cracker misses every enrolment-day cell", out["cracker"]["hit_a"] == 0, "%d of 12" % out["cracker"]["hit_a"])
    check("cracker misses every zone cell", out["cracker"]["hit_b"] == 0, "%d of 12" % out["cracker"]["hit_b"])
    ls = (out["mirror"]["hit_a"] + out["mirror"]["hit_b"]) / 24
    check("mirror keeps under two fifths of the ask weight", ls < 0.4, "%.2f" % ls)
    ga, _ = calc.gps_days(sol)
    check("figures: enrolment days", all(ga[m] == figs["ask_gps_days_2026_%02d" % (m + 1)] for m in range(12)))
    gz = calc.zones(sol)
    check("figures: zones", all(gz[m] == figs["ask_zones_2026_%02d" % (m + 1)] for m in range(12)))
    # separation: no ask device touches the call's rows
    ev = pd.DataFrame(data["gps"]["events"])
    check("separation: the call reads nothing from the Belvantis export", "gps" not in "".join(calc.BASE + ["v12"]))
    added = h[(h["supplement_no"] == "1")]
    orig = set(zip(h.loc[h["supplement_no"] == "0", "incident_number"], h.loc[h["supplement_no"] == "0", "victim_seq"]))
    only_sup = added[[(a, b) not in orig for a, b in zip(added["incident_number"], added["victim_seq"])]]
    keys = set((calc.norm_last(r.last_name), calc.norm_first(r.first_name.split(" ")[0]), r.date_of_birth)
               for r in data["epi"].itertuples())
    hits = 0
    for nm, dob in zip(only_sup["victim_name"], only_sup["victim_dob"]):
        if dob and ", " in nm:
            last, rest = nm.split(", ", 1)
            if (calc.norm_last(last), calc.norm_first(rest.split(" ")[0]), "%s-%s-%s" % (dob[6:10], dob[0:2], dob[3:5])) in keys:
                hits += 1
    check("separation: no supplement-only victim is a supervisee", hits == 0 and len(only_sup) > 0, "%d rows" % len(only_sup))
    xw = set(data["xw"]["legacy_offender_number"])
    legacy_refs = set(ev["client_ref"]) & xw
    check("ask device: some Belvantis clients carry legacy numbers", len(legacy_refs) >= 5, "%d" % len(legacy_refs))
    # Schedule C.2's same-date rule: read in logged order instead, a same-date resumption logged first leaves its
    # suspension open, so the clause decides those days; the golden pairs each suspension with the first resumption
    # effective on or after it
    evx = calc.gps_events(sol)
    same = 0
    for ref, g in evx.groupby("client_ref"):
        r_eff = set(g.loc[g["event"] == "RESUME", "eff"])
        same += sum(1 for a in g.loc[g["event"] == "SUSPEND", "eff"] if a in r_eff)
    check("Schedule C.2 files the same-date rule, and the export holds same-date suspensions it decides",
          "same date suspend that one day" in DC.BILL_SUSPEND and same > 0, "%d same-date pairs" % same)
    check("Schedule C.2 gives Merrow time's offset and daylight-saving rule",
          "six hours behind UTC" in DC.BILL_DAY and "second Sunday in March" in DC.BILL_DAY and
          data["gps"]["account_time_zone"] == "US/Central")
    # Schedule C.3: a move between Darrow and Keswick counties leaves Region 4's supervision and account in place
    sup = sol.sup
    moved = 0
    for p in set(evx["pid"]):
        eps = sup.by_person.get(p, [])
        for r0 in eps:
            if r0.close_reason in ("TRANSFER - KESWICK COUNTY", "TRANSFER - DARROW COUNTY") and r0.close is not None:
                nxt = [r1 for r1 in eps if r1.start == r0.close + dt.timedelta(days=1)]
                moved += int(bool(nxt))
    check("Schedule C.3 files the county-transfer rule, and enrolled clients move between the region's counties",
          "offices or counties leaves the account where it is" in DC.BILL_REGION and moved > 0, "%d moves" % moved)
    gen = pd.Timestamp(data["gps"]["export"]["exported_utc"])
    late = sum(1 for e in data["gps"]["events"] if pd.Timestamp(e["logged_utc"]) > gen)
    check("realism: no Belvantis event logged after the export was generated", late == 0, "%d" % late)
    marked = sum(1 for e in data["gps"]["events"] if e["event"] in ("ENROL", "END") and e["note"])
    check("ask device: no note marks an enrolment's opening or closing", marked == 0, "%d" % marked)

    # ------------------------------------------------------------ realism sweep
    print("realism sweep")
    def not_arith(vals):
        v = sorted(set(int(re.sub(r"\D", "", str(x)) or 0) for x in vals))
        d = np.diff(v)
        return len(set(d.tolist())) > 3
    check("realism: episode ids are not an arithmetic sequence", not_arith(data["epi"]["episode_id"]))
    check("realism: Darrow booking numbers are not an arithmetic sequence", not_arith(data["dj"]["booking_no"]))
    check("realism: police incident numbers are not an arithmetic sequence", not_arith(h["incident_number"].str[-6:]))
    check("realism: Belvantis event ids are not an arithmetic sequence", not_arith(ev["event_id"]))
    yearly = sol.panel.nb[sol.panel.nb["schedb"]].groupby(sol.panel.nb["date"].map(lambda d: d.year)).size()
    check("realism: no two years' Category V bookings equal", yearly.duplicated().sum() == 0, str(dict(yearly)))
    check("realism: the caseload is not a round number", len(f26) % 100 != 0, "%d" % len(f26))
    evn = pd.DataFrame(data["gps"]["events"])
    for typ in ("SUSPEND", "VOID", "ACCOUNT_TRANSFER"):
        notes = evn.loc[evn["event"] == typ, "note"].fillna("")
        check("realism: %s notes take several wordings and are sometimes blank" % typ,
              notes.nunique() >= 3 and (notes == "").any(), "%d wordings" % notes.nunique())
    ex = data["epi"][data["epi"]["close_reason"] == "EXPIRATION OF SENTENCE"]
    gap = (pd.to_datetime(ex["close_date"]) - pd.to_datetime(ex["expiration_date"])).dt.days
    check("realism: a sentence that expires closes on its expiration date or the day after",
          ((gap == 0) | (gap == 1)).all() and (gap == 0).mean() > 0.8, "%.2f the same day" % (gap == 0).mean())
    cn = data["xw"]["conversion_note"]
    check("realism: the crosswalk's conversion notes take several values", cn.nunique() >= 5, "%d values" % cn.nunique())
    biz = h.loc[h["victim_type"] == "B", "victim_name"].unique()
    shapes = set(re.sub(r"[A-Z]+", "W", re.sub(r"\d+", "9", b)) for b in biz)
    check("realism: business victims take more than two name shapes", len(shapes) > 2, "%d shapes" % len(shapes))
    people = ["Zachary Wallace", "Ashlee Evans"]
    dtext = pdf_text(os.path.join(target, C.F_DIRECTIVE))
    check("realism: the directive names the card's persona", "Ashlee Evans" in dtext)
    import subprocess
    from pdfminer.high_level import extract_text as miner_text
    for f, sentences in DC.ALL_TEXT.items():
        t = pdf_text(os.path.join(target, f))
        check("document read back: %s (positive control)" % f, len(t) > 200 and all(" ".join(s.split()) in t for s in sentences))
        # a strict reader and two independent extractors read the same filed sentences (house fix H25)
        others = [" ".join(" ".join(pg.extract_text() or "" for pg in __import__("pypdf").PdfReader(os.path.join(target, f), strict=True).pages).split()),
                  " ".join(miner_text(os.path.join(target, f)).split()),
                  " ".join(subprocess.run(["pdftotext", "-layout", os.path.join(target, f), "-"], capture_output=True,
                                          text=True).stdout.split())]
        check("document read back: %s by a strict reader, pdfminer and pdftotext" % f,
              all(len(o) > 200 and all(" ".join(s.split()) in o for s in sentences) for o in others))
    check("documents: no document states a date after its file", "2027" not in pdf_text(os.path.join(target, C.F_SCHED)))
    # house fix H16: every dated column ends by the records' end, bar an episode's expiration (a sentence's term)
    late = []
    end = pd.Timestamp("2026-12-31 23:59:59")
    pg0 = pd.read_csv(os.path.join(target, C.F_PROG), dtype=str, keep_default_na=False)
    for name, frame, cols in ((C.F_DARROW, data["dj"], ["booked", "released"]), (C.F_HPD, h, ["report_date", "occurred_date"]),
                              (C.F_EPIS, data["epi"], ["start_date", "close_date"]), (C.F_PROG, pg0, ["referral_date", "start_date", "discharge_date"]),
                              (C.F_GPS, pd.DataFrame(data["gps"]["events"]), ["effective_date"])):
        for c in cols:
            v = pd.to_datetime(frame[c].replace("", None).dropna(), format="mixed")
            if len(v) and v.max() > end:
                late.append("%s %s" % (name, c))
    kw = data["kj"]
    for c in ("Booked (UTC)", "Released (UTC)"):
        v = pd.to_datetime(kw[c].replace("", None).dropna(), utc=True).dt.tz_convert("America/Chicago").dt.tz_localize(None)
        if len(v) and v.max() > end:
            late.append("%s %s" % (C.F_KESWICK, c))
    check("realism: no dated column runs past the records' end of 31 December 2026", not late, ", ".join(late))
    for f in os.listdir(target):
        if f.endswith(".xlsx"):
            with zipfile.ZipFile(os.path.join(target, f)) as z:
                app = z.read("docProps/app.xml").decode()
            check("container: %s carries Excel's own application stamp" % f, "Microsoft Excel" in app and "openpyxl" not in app)
        if f.endswith(".pdf"):
            b = open(os.path.join(target, f), "rb").read()
            check("container: %s carries no library banner" % f, b"ReportLab" not in b)

    # ------------------------------------------------------------ the goldens read back
    print("goldens")
    import openpyxl
    from docx import Document
    gd = os.path.join(os.path.dirname(os.path.abspath(target)), "golden")
    wb = openpyxl.load_workbook(os.path.join(gd, C.DELIV_XLSX))
    days = [wb["Enrolment days 2026"].cell(row=4 + m, column=2).value for m in range(12)]
    zns = [wb["Exclusion zones 2026"].cell(row=4 + m, column=2).value for m in range(12)]
    res = [c.value for row in wb["2027 reservation"].iter_rows() for c in row]
    check("goldens: the workbook's months are the figures", days == [figs["ask_gps_days_2026_%02d" % (m + 1)] for m in range(12)]
          and zns == [figs["ask_zones_2026_%02d" % (m + 1)] for m in range(12)])
    check("goldens: the workbook reserves the call", res.count(figs["call_forecast_2027"]) >= 2)
    memo = "\n".join(p.text for p in Document(os.path.join(gd, C.DELIV_DOCX)).paragraphs)
    memo += "\n".join(c.text for t in Document(os.path.join(gd, C.DELIV_DOCX)).tables for r in t.rows for c in r.cells)
    want = ["Reserve %s enrolments" % "{:,}".format(figs["call_forecast_2027"]), "{:,}".format(figs["ask_gps_days_2026_total"]),
            "(%s in all)" % "{:,}".format(figs["ask_zones_2026_total"]), "{:,}".format(figs["cc_caseload_2026"]),
            "about %s" % "{:,}".format(figs["cc_victims_on_caseload_2026_about"]), "%.1f per cent" % figs["cc_rate_victims_pct"],
            "on its own (%s)" % "{:,}".format(figs["default_printed"])]
    miss = [w for w in want if w not in memo]
    check("goldens: the memo states the figures as the workbook and the write-up do", not miss, ", ".join(miss))

    # ------------------------------------------------------------ input gates and distractors
    print("input gates and distractors")
    files = sorted(os.listdir(target))
    check("gate: ten or more files", len(files) >= 10, "%d" % len(files))
    check("gate: three or more formats", len(set(os.path.splitext(f)[1] for f in files)) >= 3)
    big = max((sum(1 for _ in open(os.path.join(target, f), encoding="utf-8")) - 1) for f in files if f.endswith(".csv"))
    check("gate: a file of 25,000 or more rows", big >= 25000, "%d" % big)
    check("gate: no readme, dictionary, manifest or notes file", not any(re.search(r"readme|dictionar|manifest|notes|codebook|schema", f, re.I)
                                                                          for f in files))
    meta = json.load(open(os.path.join(os.path.dirname(HERE), "metadata.json")))
    check("gate: two distractors named in metadata.json", len(meta["distractor_files"]) >= 2 and
          all(os.path.exists(os.path.join(os.path.dirname(target), x)) for x in meta["distractor_files"]))
    check("gate: no file name says distractor", not any("distract" in f for f in files))
    # a solver weighing the programme referrals as a predictor still files the call
    pg = pd.read_csv(os.path.join(target, C.F_PROG), dtype=str, keep_default_na=False)
    comp = defaultdict(list)
    for p, d0, s0 in zip(pg["cms_person_id"], pg["discharge_date"], pg["discharge_status"]):
        if s0 == "COMPLETED" and d0:
            comp[p].append(D.fromisoformat(d0))
    def addc(f):
        f = f.copy()
        o = {Y: D(Y, 12, 31) for Y in set(f["Y"])}
        f["prog"] = [int(any(x <= o[Y] for x in comp.get(p, []))) for p, Y in zip(f["pid"], f["Y"])]
        return f
    trp, f26p = addc(tr), addc(f26)
    mp = calc.fit(trp, calc.BASE + ["v12", "prog"])
    vp = calc.predict(mp, f26p, calc.BASE + ["v12", "prog"])
    check("distractor: programme completion as a feature leaves the call", LO <= vp < HI, "%.1f" % vp)
    check("distractor: the recidivism brief's rates are not the call's basis", "three years" in pdf_text(os.path.join(target, C.F_BRIEF)))

    n_fail = sum(1 for _, ok, _ in RES if not ok)
    print("checks: %d assertions, %d failed" % (len(RES), n_fail))
    return 1 if n_fail else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(HERE), "target")))
