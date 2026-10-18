"""Write task170/metadata.json from the built pack: the distractors, the input gates and every shipped file's source,
date and licence."""
import json
import os

import openpyxl

import docs as DOC
import pack as PK

ORIGIN = "Constructed for this task (seeded, deterministic): "
SOURCES = {
    PK.F_LEG: "legacy paediatric ART register, visits 1 January 2023 to 30 June 2024",
    PK.F_EPR: "national EPR paediatric and adolescent ART module, encounters 1 June 2024 to 31 August 2026",
    PK.F_PTS: "national EPR paediatric and adolescent ART module, patient master at 2 September 2026",
    PK.F_CW: "national EPR migration tables, legacy ART numbers to national patient identifiers",
    PK.F_CODES: "national regimen code list, 2026",
    PK.F_PROJ: "National ART Programme patient projection, September 2026 round",
    PK.F_STOCK: "Central Medical Stores stock status report, count of 31 August 2026",
    PK.F_DEL: "Central Medical Stores warehouse system, outbound deliveries of TLD 30s, July 2025 to August 2026",
    PK.F_POD: "Central Medical Stores proof-of-delivery app, July 2025 to August 2026",
    PK.F_ACC: "Central Medical Stores customer master at 2 September 2026",
    PK.F_LMIS: "eLMIS ART programme requisitions, periods July 2025 to August 2026",
    PK.F_REG: "Ministry of Health facility registry extract, ART sites, 1 September 2026",
    PK.F_VL: "national laboratory information system, paediatric viral load results from July 2025",
    PK.F_DICT: "Central Medical Stores quantification unit, data dictionary",
    PK.F_MAN: "Central Medical Stores quantification unit, extract manifest for the 2 September 2026 pull",
    DOC.F_DOSING: "Ministry of Health, national HIV guidelines 2026, Annex 7 paediatric dosing (effective 1 March 2026)",
    DOC.F_SOP: "Central Medical Stores, SOP QP-04 version 3, cycle quantification (effective 1 March 2026)",
    DOC.F_EST: "National ART Programme, HIV estimates 2026 round, summary tables",
}
DISTRACTORS = list(PK.DISTRACTORS)
TITLE = ("How many pallets of TLD 30s a national Central Medical Stores should order on the children's line of its "
         "January to June 2027 order, when children are weighed into the 30 kg band at visits that just became less "
         "frequent")


def rows_of(path):
    if path.endswith(".csv"):
        with open(path, encoding="utf-8") as fh:
            return sum(1 for _ in fh) - 1
    if path.endswith(".xlsx"):
        wb = openpyxl.load_workbook(path, read_only=True)
        n = sum(max((ws.max_row or 1) - 1, 0) for ws in wb.worksheets)
        wb.close()
        return n
    if path.endswith(".json"):
        with open(path, encoding="utf-8") as fh:
            d = json.load(fh)
        return sum(len(q["lineItems"]) for q in d["requisitions"])
    return None


def stated_inputs(target):
    """Figures a shipped file states as they stand: the stock status report's usable TLD stock at the count."""
    wb = openpyxl.load_workbook(os.path.join(target, PK.F_STOCK), read_only=True)
    out = [int(r[3]) for r in wb["Summary"].iter_rows(min_row=5, values_only=True) if r and r[0] == "ARV0412"]
    wb.close()
    return out


def write(task_dir):
    target = os.path.join(task_dir, "target")
    out = []
    for f in sorted(os.listdir(target)):
        p = os.path.join(target, f)
        out.append({"path": f, "format": f.rsplit(".", 1)[1], "bytes": os.path.getsize(p), "rows": rows_of(p),
                    "source": ORIGIN + SOURCES[f], "date": "2026-09-02",
                    "license": "CC0-1.0 (original work, constructed for this task)"})
    big = max((x for x in out if x["rows"]), key=lambda x: x["rows"])
    meta = {
        "task": "task170",
        "title": TITLE,
        "domain": "Healthcare Delivery & Health Economics",
        "client_domain": "Healthcare Delivery & Health Economics",
        "house_domain": "healthcare-delivery-health-economics",
        "subdomain": "utilization, access and demand forecasting (a national ARV quantification: the children's line of "
                     "the C7 TLD order, the paediatric regimen transitions forecast child by child from the register's "
                     "scheduled weighings)",
        "objective": "Forecasting & Predictive Modeling",
        "prompt_shape": "the children's line of one cycle's TLD order in whole pallets with the children it carries, "
                        "the adults' line with its consumption, maximum stock and the stock and shipments it nets, each "
                        "region's CMS deliveries and reported dispensing over the last two closed cycles, and a chart of "
                        "those deliveries and that dispensing by region for each cycle, every bar labelled",
        "as_of": "2026-09-02",
        "deliverables": ["c7_tld_order.xlsx", "c4_c5_tld_flows.png"],
        "distractor_files": DISTRACTORS,
        "field_terms": [],
        "stated_inputs": stated_inputs(target),
        "input_gates": {"files": len(out), "formats": sorted({x["format"] for x in out}),
                        "largest_file": big["path"], "largest_file_rows": big["rows"]},
        "sources": [{"files": "all files under target/",
                     "origin": "constructed for this task by task170/generator/build.py (seeded, deterministic) around "
                               "a fictional southern African ministry of health, its Central Medical Stores and national "
                               "ART programme, in the layouts of a legacy paediatric register and a national EPR, a "
                               "warehouse system's delivery and customer exports, a proof-of-delivery app, an eLMIS "
                               "requisition export, a facility registry and the programme's own guideline annex, SOP, "
                               "projection and estimates; no real organisation's records or person is depicted.",
                     "date": "2026-09-02", "license": "CC0-1.0 (original work)"}],
        "created": "2026-10-18",
        "generator": "task170/generator/build.py",
        "files": out,
    }
    with open(os.path.join(task_dir, "metadata.json"), "w") as fh:
        json.dump(meta, fh, indent=2, ensure_ascii=False)
        fh.write("\n")
    return meta
