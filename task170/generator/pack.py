"""task170 pack: writes every file of the evidence pack into a folder from the clean truth (world.py, records.py)
with the Layer 1 mess each system writes. Returns row counts and the truth the checks need."""
from __future__ import annotations

import calendar
import csv
import datetime as dt
import io
import math
import os

import numpy as np
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import common as C
import docs as DOCS
import postprocess as PP
import records as R
import world as W

D = dt.date
BOLD = Font(name="Calibri", size=11, bold=True)
BODY = Font(name="Calibri", size=11)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")

F_HLC = "hlc_sentinel_2024_2026.csv"
F_VILL = "sentinel_villages.csv"
F_HBS = "human_behaviour_survey_2025.xlsx"
F_BASE = "obf_planning_baseline_2027.xlsx"
F_REG = "nets_registration_2026.parquet"
F_KHIS = "khis_moh705_obf_2024_2026.csv"
F_FAC = "facility_list_2026.csv"
F_DICT = "data_dictionary.txt"
F_MAN = "export_manifest.csv"
F_SUSC = "insecticide_susceptibility_2025.csv"
F_RAIN = "rainfall_daily_2019_2026.csv"

TRIAL_SEED = 254          # the trial's draw (see the design note's Draw)


def clock(t):
    """Hours after 18:00 to a clock time (hour, minute)."""
    mins = int(round(t * 60)) % (24 * 60)
    total = (18 * 60 + mins) % (24 * 60)
    return total // 60, total % 60


def fmt_time(t, style):
    h, m = clock(t)
    if style == "hhmm":
        return "%02d:%02d" % (h, m)
    if style == "compact":
        return "%02d%02d" % (h, m)
    if style == "ampm":
        return "%d:%02d %s" % (h % 12 or 12, m, "PM" if h >= 12 else "AM")
    if style == "bare":
        return "%d:%02d" % (h % 12 or 12, m)
    return "%02d.%02d" % (h, m)


def _csv(path, head, rows):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(head)
    for r in rows:
        w.writerow(r)
    open(path, "w", encoding="utf-8", newline="").write(buf.getvalue())


def _sheet(ws, head, rows, widths=None, start_row=1, fmt=None):
    for j, h in enumerate(head):
        c = ws.cell(row=start_row, column=j + 1, value=h)
        c.font = BOLD
        c.fill = HEAD_FILL
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = ws.cell(row=start_row + 1 + i, column=j + 1, value=v)
            c.font = BODY
            if isinstance(v, D):
                c.number_format = "dd/mm/yyyy"
            elif fmt and j in fmt:
                c.number_format = fmt[j]
    for j, h in enumerate(head):
        ws.column_dimensions[get_column_letter(j + 1)].width = (widths[j] if widths else max(10, len(h) + 2))


def _save(wb, path, creator, last_by, created, modified):
    wb.save(path)
    PP.normalize_ooxml(path, creator, last_by, created, modified, excel=True, company=C.ORG)


# ---------------------------------------------------------------------------------------------------- world
def make_world():
    S = W.trial(np.random.default_rng(TRIAL_SEED))
    hlc = R.hlc(np.random.default_rng(C.SEED + 2))
    hbs = R.survey(np.random.default_rng(C.SEED + 3))
    base = R.baseline(np.random.default_rng(C.SEED + 4))
    pools = R.name_pools(C.SEED % 100000)
    reg, rereg = R.registration(np.random.default_rng(C.SEED + 5), pools)
    facs = R.facilities(np.random.default_rng(C.SEED + 6), pools)
    kh, kh_truth = R.khis(np.random.default_rng(C.SEED + 7), facs)
    return dict(S=S, hlc=hlc, hbs=hbs, base=base, reg=reg, rereg=rereg, facs=facs, khis=kh, khis_truth=kh_truth)


# ---------------------------------------------------------------------------------------------------- writers
def write_hlc(out, Wd, rng):
    from scipy.stats import binom
    rows = []
    U = rng.random((len(Wd["hlc"]), 6))
    notes = ["light rain", "wind", "power cut nearby", "dogs barking"]
    for r, u in zip(Wd["hlc"], U):
        sc, ward, vname = R.VILLAGES[r["vi"]]
        night = r["night"]
        new_app = night >= D(2025, 7, 1)
        if new_app:
            style = "compact"
        elif sc in ("KBR", "WKL"):
            style = "ampm"
        else:
            style = "colon"
        sc_then = "KDG" if (ward == C.MOVED_WARD[0] and night < C.BOUNDARY) else sc
        pos = ("Indoor" if r["position"] == "IN" else "Outdoor") if new_app else r["position"]
        mins, nv, nt, nc, rem = r["minutes"], r["an_ruv"], r["an_ten"], r["culex"], ""
        date_s = night.strftime("%d/%m/%Y") if new_app else night.isoformat()
        if u[0] < 0.008:
            rows.append([R.village_code(r["vi"]), vname, sc_then, date_s, r["house"], pos, r["collector"],
                         C.hour_label(r["h"], style), "", "", "", "", "collector absent"])
            continue
        if u[0] < 0.028:
            mins = 30
            nv, nt, nc = (int(binom.ppf(min(max(uu, 1e-12), 1 - 1e-12), x, 30 / 45)) for uu, x in zip(u[1:4], (nv, nt, nc)))
            rem = "collector relieved"
        if rem == "" and u[4] < 0.004:
            rem = notes[int(u[5] * 4)]
        rows.append([R.village_code(r["vi"]), vname, sc_then, date_s, r["house"], pos, r["collector"],
                     C.hour_label(r["h"], style), mins, nv, nt, nc, rem])
    allrows = rows
    head = ["village_code", "village", "sub_county", "night_of", "house", "position", "collector_id", "hour",
            "minutes_collected", "an_ruvariensis", "an_tenebrosus", "culex_spp", "remarks"]
    _csv(out + "/" + F_HLC, head, allrows)
    return len(allrows), []


def write_villages(out):
    rows = []
    for i, (sc, ward, v) in enumerate(R.VILLAGES):
        rows.append([R.village_code(i), v, ward, sc, C.SC_NAME[sc], "2023-11" if i < 12 else "2024-01"])
    _csv(out + "/" + F_VILL, ["village_code", "village", "ward", "sub_county_code", "sub_county", "sentinel_since"], rows)


def write_hbs(out, Wd, rng):
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Round 1 Apr 2025"
    ws1 = wb.create_sheet("Round 2 Aug 2025")
    head = ["hh_no", "line_no", "village_code", "age_years", "sex", "went_inside", "went_to_bed", "got_up",
            "slept_under_any_net", "interview_date", "enumerator"]
    by_round = {1: [], 2: []}
    agemap = {"under 5": (0, 4), "5 to 14": (5, 14), "15 and over": (15, 84)}
    ages = {}
    for r in Wd["hbs"]:
        key = (r["hh"], r["person"])
        if key not in ages:
            lo, hi = agemap[r["group"]]
            ages[key] = (int(rng.integers(lo, hi + 1)), "F" if rng.random() < 0.52 else "M")
        age, sex = ages[key]
        if r["round"] == 2 and r.get("revisit"):
            age = age + (1 if rng.random() < 0.3 else 0)
        styles = []
        for t in (r["t_in"], r["t_bed"], r["t_rise"]):
            u = rng.random()
            st = "hhmm" if u < 0.72 else ("ampm" if u < 0.90 else "compact")
            styles.append(fmt_time(t, st))
        d0 = D(2025, 4, 7) if r["round"] == 1 else D(2025, 8, 4)
        day = d0 + dt.timedelta(days=int(rng.integers(0, 19)))
        net = "Yes" if rng.random() < 0.14 else "No"
        hh_no = "HB25-%s-%04d" % (R.village_code(r["vi"])[3:6], 1000 + r["hh"])
        by_round[r["round"]].append([hh_no, r["person"], R.village_code(r["vi"]), age, sex] + styles +
                                    [net, day, "EN%02d" % (1 + (r["hh"] * 7) % 23)])
    for ws, k in ((ws0, 1), (ws1, 2)):
        rows = sorted(by_round[k], key=lambda x: (x[0], x[1]))
        _sheet(ws, head, rows)
    _save(wb, out + "/" + F_HBS, C.PEOPLE["ento"], C.PEOPLE["ento"], "2025-04-03T07:12:00Z", "2025-09-02T13:47:00Z")
    return len(by_round[1]) + len(by_round[2])


def write_baseline(out, Wd):
    wb = Workbook()
    ws = wb.active
    ws.title = "Baseline 2027"
    ws.cell(row=1, column=1, value="Ruvari County Ombeni fever planning baseline, 2027: expected confirmed cases by "
                                   "sub-county and month without the LulaNet R campaign. County planning model run of "
                                   "14 September 2026.").font = BODY
    rows = []
    for sc in C.SC_CODES:
        for (y, m) in C.WINDOW:
            rows.append([sc, C.SC_NAME[sc], "%s %d" % (calendar.month_abbr[m], y), Wd["base"][(sc, m)]])
    _sheet(ws, ["sub_county_code", "sub_county", "month", "expected_cases"], rows, widths=[16, 16, 12, 16], start_row=3)
    _save(wb, out + "/" + F_BASE, C.PEOPLE["vbd"], C.PEOPLE["vbd"], "2026-09-14T08:03:00Z", "2026-09-14T08:41:00Z")


def write_registration(out, Wd):
    reg = Wd["reg"]
    cols = {
        "reg_no": pa.array([r["reg_no"] for r in reg], pa.string()),
        "round": pa.array([r["round"] for r in reg], pa.string()),
        "reg_date": pa.array([r["reg_date"] for r in reg], pa.date32()),
        "chv_code": pa.array([r["chv"] for r in reg], pa.string()),
        "sub_county": pa.array([r["sub_county"] for r in reg], pa.string()),
        "ward": pa.array([r["ward"] for r in reg], pa.string()),
        "village": pa.array([r["village"] for r in reg], pa.string()),
        "dwelling_mark": pa.array([r["dwelling"] for r in reg], pa.string()),
        "head_name": pa.array([r["head"] for r in reg], pa.string()),
        "head_sex": pa.array([r["sex"] for r in reg], pa.string()),
        "head_id_no": pa.array([r["id_no"] or None for r in reg], pa.string()),
        "phone": pa.array([r["phone"] or None for r in reg], pa.string()),
        "usual_members": pa.array([r["members"] for r in reg], pa.int16()),
        "under_five": pa.array([r["under5"] for r in reg], pa.int16()),
        "sleeping_spaces": pa.array([r["spaces"] for r in reg], pa.int16()),
        "remarks": pa.array([r["remarks"] or None for r in reg], pa.string()),
    }
    tb = pa.table(cols)
    pq.write_table(tb, out + "/" + F_REG, compression="snappy", use_dictionary=True, write_statistics=False,
                   store_schema=False)
    return len(reg)


def write_facilities(out, Wd):
    rows = []
    for f in Wd["facs"]:
        rows.append([f["mfl"], f["name"], f["level"], C.SC_NAME[f["sub_county"]], f["ward"], f["opened"].isoformat(),
                     f["reporting_since"].strftime("%Y-%m"), f["parent"], "Operational"])
    _csv(out + "/" + F_FAC, ["mfl_code", "facility_name", "keph_level", "sub_county", "ward", "opened",
                             "reporting_since", "parent_mfl_code", "status"], rows)


def write_khis(out, Wd, rng):
    facs = {f["mfl"]: f for f in Wd["facs"]}
    rows = []
    for r in Wd["khis"]:
        f = facs[r["mfl"]]
        sc = f["sub_county"]
        moved = f["ward"] == C.MOVED_WARD[0]
        then = "KDG" if (moved and (r["y"], r["m"]) < (2026, 7)) else sc
        path = "Kenya/Ruvari County/%s Sub County/%s Ward/%s" % (C.SC_NAME[then], f["ward"], f["name"])
        nxt = (r["y"] + (r["m"] == 12), r["m"] % 12 + 1)
        upd = D(nxt[0], nxt[1], int(rng.integers(3, 12)))
        age = "Under 5 years" if r["age"] == "<5" else "5 years and over"
        rows.append(["%04d%02d" % (r["y"], r["m"]), path, r["mfl"], f["name"], C.SC_NAME[then], age, r["tested"],
                     r["confirmed"], r["clinical"], upd.isoformat(), r["comment"]])
        if moved and (r["y"], r["m"]) == (2026, 6):
            path2 = "Kenya/Ruvari County/%s Sub County/%s Ward/%s" % (C.SC_NAME[sc], f["ward"], f["name"])
            rows.append(["%04d%02d" % (r["y"], r["m"]), path2, r["mfl"], f["name"], C.SC_NAME[sc], age, r["tested"],
                         r["confirmed"], r["clinical"], D(2026, 7, 21).isoformat(), ""])
    rows.sort(key=lambda x: (x[0], x[1], x[5]))
    _csv(out + "/" + F_KHIS, ["period", "orgunit_path", "mfl_code", "facility", "sub_county", "age_group",
                              "tested_obf", "confirmed_obf", "clinical_obf", "last_updated", "comment"], rows)
    return len(rows)


def write_susceptibility(out, rng):
    rows = []
    for i, (sc, ward, v) in enumerate(R.VILLAGES[:8]):
        for ins, conc in (("Permethrin", "0.75%"), ("Deltamethrin", "0.05%"), ("Alpha-cypermethrin", "0.05%"),
                          ("Pirimiphos-methyl", "0.25%")):
            n = int(rng.integers(92, 104))
            p = {"Permethrin": 0.71, "Deltamethrin": 0.78, "Alpha-cypermethrin": 0.74, "Pirimiphos-methyl": 0.99}[ins]
            dead = int(rng.binomial(n, min(0.999, p * float(np.exp(rng.normal(0, 0.06))))))
            rows.append([R.village_code(i), v, D(2025, 5 + i % 3, int(rng.integers(2, 26))).isoformat(),
                         "An. ruvariensis", ins, conc, n, dead, "%.1f" % (100 * dead / n)])
    _csv(out + "/" + F_SUSC, ["village_code", "village", "test_date", "species", "insecticide", "concentration",
                              "exposed", "dead_24h", "mortality_pct"], rows)


def write_rain(out, rng):
    rows = []
    stations = [("RVR-MET-01", "Ruvari Town"), ("RVR-MET-02", "Kobura"), ("RVR-MET-03", "Wakula")]
    d = D(2019, 1, 1)
    while d <= D(2026, 9, 30):
        for code, name in stations:
            wet = R.SEASON[d.month]
            mm = 0.0 if rng.random() > 0.18 * wet + 0.05 else round(float(rng.gamma(1.1, 6.5 * wet)), 1)
            rows.append([code, name, d.isoformat(), "%.1f" % mm if rng.random() > 0.004 else ""])
        d += dt.timedelta(days=1)
    _csv(out + "/" + F_RAIN, ["station_code", "station", "date", "rain_mm"], rows)


DICTIONARY = """Ruvari County Ombeni fever programme: data dictionary for the files pulled on 12 October 2026
Prepared by Tulivu Health Partners monitoring and evaluation with the county entomology unit.

hlc_sentinel_2024_2026.csv
  One row per collector per hour of a catch night: two houses a night, one collector indoors and one outdoors at each
  house, at the county's sentinel villages, January 2024 to September 2026.
  village_code, village     the sentinel village (sentinel_villages.csv)
  sub_county                the sub-county the village sat in on the night of the catch
  night_of                  the date of the evening the catch began (ISO dates to June 2025; dd/mm/yyyy from the
                            tablet app, July 2025 on)
  house                     1 or 2
  position                  IN or OUT to June 2025; Indoor or Outdoor from July 2025
  collector_id              the collector
  hour                      the hour of collection, 18:00 to 06:00; the label format follows the team's sheet
  minutes_collected         minutes catching in the hour (45 as standard)
  an_ruvariensis            Anopheles ruvariensis caught landing (the county's Ombeni fever vector)
  an_tenebrosus             Anopheles tenebrosus caught landing
  culex_spp                 culicines caught landing
  remarks                   the collector's note

sentinel_villages.csv
  One row per sentinel village, with its ward and sub-county under the July 2026 realignment.

human_behaviour_survey_2025.xlsx
  One row per household member interviewed, two rounds (sheets): April 2025 and August 2025. Round 2 revisited some
  round 1 households and added others. Every member of a sampled household is interviewed (a carer answers for a
  child).
  went_inside               the time the member went indoors for the night, the evening before the interview
  went_to_bed               the time the member went to bed that night
  got_up                    the time the member got up the next morning
  slept_under_any_net       the member slept under a net of any kind that night
  Times are as the enumerator wrote them.

obf_planning_baseline_2027.xlsx
  Expected confirmed Ombeni fever cases by sub-county and month, April to December 2027, without the LulaNet R
  campaign: the county planning model's run of 14 September 2026.

nets_registration_2026.parquet
  One row per household registration made by a community health volunteer in the campaign registration: the main
  round (7 to 26 September 2026) and the mop-up (1 to 9 October 2026).
  reg_no                    the registration number written on the household's slip
  round                     MAIN or MOPUP
  dwelling_mark             the chalk mark the registration team writes on the dwelling's door
  head_name, head_sex       the person registered as household head
  head_id_no, phone         as given; blank when not given
  usual_members             usual members of the household
  under_five                usual members under five years
  sleeping_spaces           places members sleep (beds or mats)

khis_moh705_obf_2024_2026.csv
  Monthly outpatient returns (MOH 705) for Ombeni fever from the county's KHIS export, one row per facility, month and
  age group: tested, confirmed and clinically diagnosed. sub_county and orgunit_path are the reporting hierarchy the
  return was entered under. last_updated is the return's last edit.

facility_list_2026.csv
  The county's master facility list at the pull, sub-counties as realigned on 1 July 2026. parent_mfl_code is the
  facility that supervises an outreach clinic. reporting_since is the first month the facility reported under its own
  code.

insecticide_susceptibility_2025.csv
  WHO tube tests on An. ruvariensis collected as larvae at eight sentinel villages, 2025.

rainfall_daily_2019_2026.csv
  Daily rainfall at the county's three weather stations.

kinga5_final_report_2024.pdf
  The KINGA-5 trial's final report.

lulanet_r_distribution_plan_2027.docx
  The campaign's operational plan.

export_manifest.csv
  What each file in this pull is, who owns it and the period it covers.
"""


def write_dictionary(out):
    open(out + "/" + F_DICT, "w", encoding="utf-8", newline="").write(DICTIONARY.replace("\n", "\r\n"))


MANIFEST = [
    (F_HLC, "County entomology unit", "Sentinel human landing catch sheets, keyed from paper to June 2025 and from the "
     "tablet app from July 2025", "2024-01 to 2026-09"),
    (F_VILL, "County entomology unit", "Sentinel village register", "as at 2026-10-12"),
    (F_HBS, "County entomology unit", "Human behaviour survey, two rounds", "2025-04 and 2025-08"),
    (F_BASE, "County Department of Health", "Planning model run 2026-09-14", "2027-04 to 2027-12"),
    (F_REG, "Tulivu Health Partners", "Campaign registration app export", "2026-09-07 to 2026-10-09"),
    (F_KHIS, "County health records office", "KHIS analytics export, MOH 705 Ombeni fever data elements", "2024-01 to 2026-09"),
    (F_FAC, "County health records office", "Master facility list", "as at 2026-10-12"),
    (F_SUSC, "County entomology unit", "Insecticide susceptibility tests", "2025"),
    (F_RAIN, "County meteorological office", "Daily station rainfall", "2019-01-01 to 2026-09-30"),
    (DOCS.F_TRIAL, "KINGA-5 Trial Steering Committee", "Trial final report (shared with the trial's permission)", "2022-02 to 2023-01"),
    (DOCS.F_PLAN, "County Department of Health and Tulivu Health Partners", "Operational plan v2.1", "2027-03 distribution"),
    (F_DICT, "Tulivu Health Partners", "Field definitions for the files in this pull", "as at 2026-10-12"),
]


def write_manifest(out):
    _csv(out + "/" + F_MAN, ["file", "owner", "what_it_is", "covers"], [list(r) for r in MANIFEST])


def build_pack(out):
    os.makedirs(out, exist_ok=True)
    for f in os.listdir(out):
        os.remove(os.path.join(out, f))
    Wd = make_world()
    rng = np.random.default_rng(C.SEED + 50)
    n = {}
    n[F_HLC], Wd["hlc_dup"] = write_hlc(out, Wd, np.random.default_rng(C.SEED + 51))
    write_villages(out)
    n[F_HBS] = write_hbs(out, Wd, np.random.default_rng(C.SEED + 52))
    write_baseline(out, Wd)
    n[F_REG] = write_registration(out, Wd)
    write_facilities(out, Wd)
    n[F_KHIS] = write_khis(out, Wd, np.random.default_rng(C.SEED + 53))
    write_susceptibility(out, np.random.default_rng(C.SEED + 54))
    write_rain(out, np.random.default_rng(C.SEED + 55))
    Wd["trial_pdf"] = DOCS.write_trial(out, Wd["S"])
    DOCS.write_plan(out, R.USAGE)
    write_dictionary(out)
    write_manifest(out)
    return n, Wd
