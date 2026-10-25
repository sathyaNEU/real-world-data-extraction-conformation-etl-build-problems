"""One-sheet and multi-sheet workbooks written with openpyxl and saved in Excel's own dialect (postprocess.py)."""
import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import postprocess as PP

STAMPS = {}


def _sheet(ws, head, rows, widths=None, money_cols=(), title=None, date_fmt="m/d/yyyy"):
    r0 = 1
    if title:
        ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=12)
        r0 = 3
    for j, h in enumerate(head, 1):
        c = ws.cell(row=r0, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
    n = 0
    for i, r in enumerate(rows, r0 + 1):
        for j, v in enumerate(r, 1):
            if v is None or v == "":
                continue
            c = ws.cell(row=i, column=j, value=v)
            if isinstance(v, dt.date):
                c.number_format = date_fmt
            elif (j - 1) in money_cols:
                c.number_format = "#,##0.00"
        n += 1
    if widths:
        from openpyxl.utils import get_column_letter
        for j, wv in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = wv
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    return n


def write_table(path, sheet, head, rows, widths=None, title=None, money_cols=(), creator="Corran Vale Railways", modified=None,
                date_fmt="dd/mm/yyyy"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    n = _sheet(ws, head, rows, widths, money_cols, title, date_fmt=date_fmt)
    wb.save(path)
    who, when = STAMPS.get(path.rsplit("/", 1)[-1], ("Corran Vale Railways", "2026-10-14T03:10:00Z"))
    PP.normalize_ooxml(path, who, who, when, when, excel=True)
    return n


def write_workbook(path, sheets, widths=None, creator="Corran Vale Railways", date_fmt="dd/mm/yyyy"):
    """sheets: [(title, head, rows)]; every sheet frozen under its header."""
    wb = Workbook()
    first = True
    n = 0
    for title, head, rows in sheets:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = title
        n += _sheet(ws, head, rows, widths, date_fmt=date_fmt)
    wb.save(path)
    who, when = STAMPS.get(path.rsplit("/", 1)[-1], ("Corran Vale Railways", "2026-10-14T03:10:00Z"))
    PP.normalize_ooxml(path, who, who, when, when, excel=True)
    return n
