"""task170 goldens: the forecast workbook and the depot readiness chart, figures.json, submission.md and
metadata.json, every figure computed once here from the pack."""
from __future__ import annotations

import json
import os

import numpy as np

import common as C
import calc

NAME = {d[0]: d[1] for d in C.DEPOTS}


def r250(x):
    return int(250 * round(x / 250.0))


def r1000(x):
    return int(1000 * round(x / 1000.0))


def compute(target):
    r = calc.compute(target)
    o = dict(r)
    o["call_r"] = r250(r["call"])
    o["starts_r"] = r1000(r["starts_fwd"])
    o["n_starts"] = int(len(r["starts"]))
    o["n_errors"] = int(len(r["errors"]))
    o["n_prior"] = int((r["p4"]["prior"] > 0).sum())
    o["first_band"] = float(r["rates"][0])
    o["last_band"] = float(r["rates"][-1])
    return o


def figures(o):
    F = {"call_tde_errors_phase4_first_12_months": o["call_r"],
         "_unrounded": {"call_tde_errors_phase4_first_12_months": round(o["call"], 1),
                        "cc2_phase4_etcs_starts_first_12_months": round(o["starts_fwd"], 1)},
         "cc1_phase4_drivers": o["drivers"],
         "cc2_phase4_etcs_starts_first_12_months": o["starts_r"]}
    for c in C.P4_CODES:
        a, n = o["fleet"][c]
        F["ask_units_accepted_by_golive_%s" % c.lower()] = a
        F["ask_units_not_accepted_by_golive_%s" % c.lower()] = n
    for c in C.REC_CODES:
        l, b = o["events"][c]
        F["ask_level_transition_errors_2026_%s" % c.lower()] = l
        F["ask_brake_interventions_2026_%s" % c.lower()] = b
    F["_path"] = {"record_driver_starts": o["n_starts"], "tde_errors_recorded": o["n_errors"],
                  "phase4_drivers_with_etcs_starts": o["n_prior"],
                  "error_rate_first_50_starts": round(o["first_band"], 4),
                  "error_rate_beyond_600_starts": round(o["last_band"], 4)}
    F["_stated"] = {"rounding_of_the_call": 250, "rounding_of_the_starts": 1000, "band_width": 50, "band_top": 600,
                    "fleet_class": 461}
    return F


# ------------------------------------------------------------------ the workbook
def write_workbook(path, o):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="D9E1F2")
    wb = Workbook()
    ws = wb.active
    ws.title = "Forecast"
    ws["A1"] = ("About %s train data entry errors expected in the six Phase 4 depots' first twelve months on ETCS"
                % "{:,}".format(o["call_r"]))
    ws["A1"].font = Font(bold=True, size=12)
    rows = [("Train data entry errors the safety desk will record (expected, nearest 250)", o["call_r"]),
            ("Drivers on the six Phase 4 rosters at 31 October 2026", o["drivers"]),
            ("Their ETCS Starts of Mission in the twelve months (nearest 1,000)", o["starts_r"])]
    for i, (k, v) in enumerate(rows, 3):
        ws.cell(row=i, column=1, value=k)
        c = ws.cell(row=i, column=2, value=v)
        c.number_format = "#,##0"
        c.font = bold
    ws.cell(row=8, column=1, value="By depot (working figures; the forecast is their total rounded once)").font = bold
    head = ["Depot", "ETCS go-live", "Drivers", "Starts in the twelve months", "Expected errors"]
    for j, h in enumerate(head, 1):
        c = ws.cell(row=9, column=j, value=h)
        c.font = bold
        c.fill = fill
    p4 = o["p4"]
    for i, c in enumerate(C.P4_CODES, 10):
        sub = p4[p4["Depot"] == c]
        vals = [NAME[c], C.DEPOT[c][3], int(len(sub)), int(round(float(sub["starts"].sum()), -2)),
                int(round(o["by_depot"][c]))]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=j, value=v)
            if j == 2:
                cell.number_format = "dd/mm/yyyy"
            elif j > 2:
                cell.number_format = "#,##0"
    ws.column_dimensions["A"].width = 74
    for col in "BCDE":
        ws.column_dimensions[col].width = 16
    # the curve the forecast runs along
    wc = wb.create_sheet("Error rate by start")
    for j, h in enumerate(["Driver's ETCS starts before this one", "Starts in the record", "Errors", "Errors per start"], 1):
        c = wc.cell(row=1, column=j, value=h)
        c.font = bold
        c.fill = fill
    T, E = o["T"], o["E"]
    edges = calc.BANDS
    k = np.arange(len(T))
    b = np.digitize(k, edges) - 1
    for i in range(len(edges) - 1):
        lab = "%d to %d" % (edges[i], edges[i + 1] - 1) if edges[i + 1] < 10 ** 8 else "%d and more" % edges[i]
        st, er = float(T[b == i].sum()), float(E[b == i].sum())
        wc.append([lab, int(st), int(er), round(er / st, 4)])
        wc.cell(row=wc.max_row, column=2).number_format = "#,##0"
        wc.cell(row=wc.max_row, column=3).number_format = "#,##0"
        wc.cell(row=wc.max_row, column=4).number_format = "0.0000"
    wc.column_dimensions["A"].width = 36
    for col in "BCD":
        wc.column_dimensions[col].width = 18
    wc.freeze_panes = "A2"
    # ask A
    wa = wb.create_sheet("Fleet readiness")
    for j, h in enumerate(["Depot", "ETCS go-live", "Units allocated for 2027", "Accepted for ETCS service by go-live",
                           "Not accepted by go-live"], 1):
        c = wa.cell(row=1, column=j, value=h)
        c.font = bold
        c.fill = fill
    for c in C.P4_CODES:
        a, n = o["fleet"][c]
        wa.append([NAME[c], C.DEPOT[c][3], a + n, a, n])
        wa.cell(row=wa.max_row, column=2).number_format = "dd/mm/yyyy"
    for col, wd in zip("ABCDE", (16, 14, 22, 34, 22)):
        wa.column_dimensions[col].width = wd
    wa.freeze_panes = "A2"
    # ask C
    we = wb.create_sheet("Phase 1-3 events 2026")
    for j, h in enumerate(["Depot", "Level transition errors", "ETCS brake interventions"], 1):
        c = we.cell(row=1, column=j, value=h)
        c.font = bold
        c.fill = fill
    for c in C.REC_CODES:
        l, bk = o["events"][c]
        we.append([NAME[c], l, bk])
    for col, wd in zip("ABC", (20, 24, 26)):
        we.column_dimensions[col].width = wd
    we.freeze_panes = "A2"
    wn = wb.create_sheet("Notes")
    notes = [
        "Forecast: the chance of a train data entry error at a Start of Mission, read from the Phase 1 to 3 record "
        "in bands of 50 on each driver's own count of ETCS starts (to 600, then one band), and each driver on a "
        "Phase 4 roster run along it from the starts the driver has already made, for twelve months at the driver's "
        "2026 monthly cab activations (duties worked, latest issue). Expected number, nearest 250.",
        "Errors: register category TDE, withdrawn reports out, a linked report counted once; each error placed at its "
        "start by SoM Ref, else at the driver's last start before the report. Test starts by fleet technicians are not "
        "driver starts. Staff numbers chained through the Resourcing crosswalk.",
        "Fleet readiness: each allocated unit's two driving vehicles as the 31 October 2026 formation list holds them; a "
        "unit is accepted for ETCS service by go-live when both vehicles' accepted or planned acceptance dates fall on "
        "or before the depot's go-live date. The works records carry the fitment programme's own unit numbers.",
        "Phase 1 to 3 events: register categories LTR and BRK, January to October 2026, withdrawn reports out and linked "
        "reports once; a recorder download report (recorder time, UTC) is the same event as a driver's or signaller's "
        "report of the same driver, unit and category up to 30 minutes after it.",
        "Sources: on-board recorder downloads, safety desk register, Resourcing roster and crosswalk, crew management duty "
        "records, ETCS rollout plan, Fitment Programme Office works records, Fleet Engineering formations and "
        "renumbering register, fleet allocation 2027.",
    ]
    wn["A1"] = "Notes"
    wn["A1"].font = bold
    for i, t in enumerate(notes, 3):
        wn.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True, vertical="top")
    wn.column_dimensions["A"].width = 110
    wb.save(path)
    import postprocess as PP
    PP.normalize_ooxml(path, "Kate Potter", "Kate Potter", "2026-11-06T09:12:00Z", "2026-11-09T14:48:00Z", excel=True)


# ------------------------------------------------------------------ the chart, written as SVG text
def write_chart(path, o):
    W, H = 780, 450
    left, right, top, bottom = 70, 30, 92, 82
    acc_c, not_c, ink, muted, grid = "#2a78d6", "#eb6834", "#1b1b1b", "#5f5f5c", "#e6e5e1"
    vals = [(NAME[c], o["fleet"][c][0], o["fleet"][c][1]) for c in C.P4_CODES]
    tot_a = sum(a for _, a, _ in vals)
    tot = sum(a + n for _, a, n in vals)
    ymax = max(a + n for _, a, n in vals)
    ytop = int(np.ceil(ymax / 2.0) * 2)
    ph = H - top - bottom
    pw = W - left - right
    bw = pw / len(vals) * 0.52
    sx = lambda i: left + pw * (i + 0.5) / len(vals)
    sy = lambda v: top + ph * (1 - v / float(ytop))
    out = ['<svg xmlns="http://www.w3.org/2000/svg" width="%d" height="%d" viewBox="0 0 %d %d" '
           'font-family="Arial, Helvetica, sans-serif">' % (W, H, W, H),
           '<rect width="%d" height="%d" fill="#fcfcfb"/>' % (W, H),
           '<text x="%d" y="30" font-size="15" font-weight="bold" fill="%s">%d of the %d Class 461 units will be accepted '
           'for ETCS service by their depot\'s go-live</text>' % (left, ink, tot_a, tot),
           '<text x="%d" y="50" font-size="11" fill="%s">Units allocated for the 2027 timetable, on the fitment programme\'s '
           'current dates; a unit is accepted when both driving vehicles are</text>' % (left, muted)]
    for v in range(0, ytop + 1, 2):
        y = sy(v)
        out.append('<line x1="%d" y1="%.1f" x2="%d" y2="%.1f" stroke="%s" stroke-width="1"/>' % (left, y, W - right, y, grid))
        out.append('<text x="%d" y="%.1f" font-size="10" fill="%s" text-anchor="end">%d</text>' % (left - 8, y + 3.5, muted, v))
    for i, (nm, a, n) in enumerate(vals):
        x = sx(i) - bw / 2
        y0, y1, y2 = sy(0), sy(a), sy(a + n)
        gl = C.DEPOT[C.P4_CODES[i]][3]
        if a:
            out.append('<rect x="%.1f" y="%.1f" width="%.1f" height="%.1f" fill="%s"><title>%s: %d accepted by go-live '
                       '(%d %s)</title></rect>' % (x, y1, bw, y0 - y1, acc_c, nm, a, gl.day, gl.strftime("%B %Y")))
            out.append('<text x="%.1f" y="%.1f" font-size="12" font-weight="bold" fill="#ffffff" text-anchor="middle">%d</text>'
                       % (sx(i), (y0 + y1) / 2 + 4, a))
        if n:
            out.append('<rect x="%.1f" y="%.1f" width="%.1f" height="%.1f" fill="%s"><title>%s: %d not accepted by go-live'
                       '</title></rect>' % (x, y2, bw, max(0.0, y1 - y2 - 2), not_c, nm, n))
            out.append('<text x="%.1f" y="%.1f" font-size="12" font-weight="bold" fill="%s" text-anchor="middle">%d</text>'
                       % (sx(i), (y1 + y2) / 2 + 3, ink, n))
        out.append('<text x="%.1f" y="%.1f" font-size="11.5" fill="%s" text-anchor="middle">%s</text>' % (sx(i), y0 + 18, ink, nm))
        out.append('<text x="%.1f" y="%.1f" font-size="9.5" fill="%s" text-anchor="middle">go-live %d %s</text>'
                   % (sx(i), y0 + 32, muted, gl.day, gl.strftime("%b %Y")))
    out.append('<line x1="%d" y1="%.1f" x2="%d" y2="%.1f" stroke="%s" stroke-width="1"/>' % (left, sy(0), W - right, sy(0), ink))
    out.append('<text x="18" y="%.1f" font-size="11" fill="%s" transform="rotate(-90 18 %.1f)" '
               'text-anchor="middle">Units</text>' % (top + ph / 2, muted, top + ph / 2))
    lx = left
    out.append('<rect x="%d" y="64" width="12" height="12" fill="%s"/>' % (lx, acc_c))
    out.append('<text x="%d" y="74" font-size="11" fill="%s">Accepted by go-live</text>' % (lx + 17, ink))
    out.append('<rect x="%d" y="64" width="12" height="12" fill="%s"/>' % (lx + 150, not_c))
    out.append('<text x="%d" y="74" font-size="11" fill="%s">Not accepted by go-live</text>' % (lx + 167, ink))
    out.append('<text x="%d" y="%d" font-size="9.5" fill="%s">Source: Fitment Programme Office works records; Fleet '
               'Engineering unit formations at 31 October 2026; fleet allocation 2027.</text>' % (left, H - 14, muted))
    out.append("</svg>")
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(out) + "\n")


# ------------------------------------------------------------------ submission.md
def write_submission(path, o, F):
    L = []
    c = lambda x: "{:,}".format(x)
    L.append("## Tags\n")
    L.append("**Domain:** Human factors and ergonomics, human reliability assessment.  ")
    L.append("**Analytical objective:** Forecasting and predictive modeling.\n")
    L.append("## 1. Final Recommendation\n")
    L.append("**Plan the Phase 4 safety case and its debriefing for about %s train data entry errors recorded against "
             "the six Phase 4 depots' drivers in each depot's first twelve months on ETCS, to the nearest 250.**"
             % c(o["call_r"]))
    L.append("The figure reads the chance of an error at each Start of Mission at the driver's own count of ETCS starts "
             "in the Phase 1 to 3 record, and runs each of the %s Phase 4 drivers along it at that driver's 2026 cab "
             "activations, about %s starts in the twelve months. Not the earlier depots' first-year errors per start, "
             "nor a curve on months since go-live: both carry the suburban drivers' pace, and the rural diagrams keep a "
             "Phase 4 driver in the record's first weeks for most of the year." % (c(o["drivers"]), c(o["starts_r"])))
    L.append("")
    L.append("## 2. Critical Components\n")
    L.append("1. Drivers on the six Phase 4 depots' rosters at 31 October 2026: **%s**" % c(o["drivers"]))
    L.append("2. Their ETCS Starts of Mission in the depots' first twelve months: **%s** (nearest 1,000)" % c(o["starts_r"]))
    L.append("")
    L.append("## 3. Step-by-Step Solution\n")
    L.append("1. Placed every Phase 1 to 3 driver start in %s on the driver's own count of ETCS starts, staff numbers "
             "chained through %s and TST test starts left out: %s starts." % (C.F_SOM, C.F_XWALK, c(o["n_starts"])))
    L.append("2. Took the train data entry errors from %s (category TDE, withdrawn reports out, a linked report once): "
             "%s, each placed at its start by the SoM Ref of its own or a linked report, else at the driver's last start "
             "before the report (the register in local time, the recorder in UTC)." % (C.F_REG, c(o["n_errors"])))
    L.append("3. Read the errors per start in bands of 50 on that count, to 600 and then one band: %s in a driver's first "
             "50 starts and %s beyond 600." % ("{:.4f}".format(o["first_band"]), "{:.4f}".format(o["last_band"])))
    L.append("4. For each of the %s drivers on a Phase 4 roster in %s, took the ETCS starts already made (%s drivers "
             "transferred in with some) and twelve months at the driver's monthly cab activations on duties worked in "
             "%s, latest issue: %s starts." % (c(o["drivers"]), C.F_ROSTER, o["n_prior"], C.F_DUTY,
                                                "{:,.1f}".format(o["starts_fwd"])))
    L.append("5. Summed each driver's chance along the bands from the starts already made: %s, about %s to the nearest "
             "250." % ("{:,.1f}".format(o["call"]), c(o["call_r"])))
    L.append("6. Units: each unit in %s by its two driving vehicles in %s, accepted for ETCS service by its depot's "
             "go-live in the rollout plan when both vehicles' Accepted or Planned Acceptance dates in %s fall on or "
             "before it; the works records' Unit is the programme's own unit, so vehicles reformed into other units and "
             "renumbered units are read through the formation list." % (C.F_ALLOC, C.F_FORM, C.F_WORKS))
    L.append("7. Events: the register's LTR and BRK reports received January to October 2026, withdrawn reports out and "
             "linked reports once, each recorder download report (stamped in the recorder's UTC) folded into a driver's or "
             "signaller's report of the same driver, unit and category up to 30 minutes after it, and a later event the "
             "same day kept as its own.")
    L.append("")
    L.append("## 4. Deliverable Answers\n")
    L.append("### %s\n" % C.DELIV_XLSX)
    L.append("1. Train data entry errors the safety desk will record against the six Phase 4 depots' drivers in each "
             "depot's first twelve months on ETCS: **%s** (nearest 250)" % c(o["call_r"]))
    L.append("2. Drivers covered: **%s**; their ETCS starts over those twelve months: **%s** (nearest 1,000)"
             % (c(o["drivers"]), c(o["starts_r"])))
    L.append("3. Class 461 units allocated for the 2027 timetable, accepted for ETCS service by the depot's go-live and "
             "not accepted:")
    for d in C.P4_CODES:
        a, n = o["fleet"][d]
        L.append("   - %s: **%d** accepted, **%d** not accepted" % (NAME[d], a, n))
    L.append("4. Level transition errors and ETCS brake interventions of each Phase 1 to 3 depot's drivers, January to "
             "October 2026:")
    for d in C.REC_CODES:
        lt, bk = o["events"][d]
        L.append("   - %s: **%d** level transition errors, **%d** brake interventions" % (NAME[d], lt, bk))
    L.append("")
    L.append("### %s\n" % C.DELIV_SVG)
    L.append("1. One stacked bar per Phase 4 depot, its accepted and not accepted units, each segment labelled with its "
             "count as in answer 3 above.")
    open(path, "w", encoding="utf-8").write("\n".join(L) + "\n")


def write_metadata(path):
    files = [C.F_SOM, C.F_REG, C.F_ROSTER, C.F_XWALK, C.F_DUTY, C.F_ROLL, C.F_CONV, C.F_GUIDE, C.F_ALLOC, C.F_WORKS,
             C.F_FORM, C.F_RENUM, C.F_ACCEPT, C.F_DICT, C.F_PUNC, C.F_WX]
    meta = {"task": C.TASK,
            "distractor_files": ["target/" + f for f in C.DISTRACTORS],
            "deliverables": [C.DELIV_XLSX, C.DELIV_SVG],
            "objective": "Forecasting & Predictive Modeling",
            "domain": "Human Factors & Ergonomics / Human Reliability Assessment",
            "setting": "Arnhallow, Corran Vale Railways, its depots and units and the Arnhallow Office of Rail Safety are "
                       "fictional; every record is synthetic, built by the task's generator, and no real operator, "
                       "regulator, system or person is represented. ETCS is named as the signalling standard.",
            "sources": {f: {"source": "synthetic (task generator)", "date": C.ASSEMBLED.isoformat(), "license": "CC0-1.0"}
                        for f in files}}
    json.dump(meta, open(path, "w"), indent=2)
    open(path, "a").write("\n")


def main(target, golden, figpath, subpath, metapath, quiet=False):
    os.makedirs(golden, exist_ok=True)
    o = compute(target)
    F = figures(o)
    json.dump(F, open(figpath, "w"), indent=2, sort_keys=True)
    open(figpath, "a").write("\n")
    write_workbook(os.path.join(golden, C.DELIV_XLSX), o)
    write_chart(os.path.join(golden, C.DELIV_SVG), o)
    write_submission(subpath, o, F)
    write_metadata(metapath)
    if not quiet:
        print("call %.1f -> %d; drivers %d; starts %.1f -> %d" % (o["call"], o["call_r"], o["drivers"], o["starts_fwd"],
                                                                  o["starts_r"]))
        print("fleet", o["fleet"])
        print("events", o["events"])
