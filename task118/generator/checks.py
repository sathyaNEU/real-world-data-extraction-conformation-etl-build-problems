"""task118 checks, run on every build: the three rules (R1 to R3c), the ladder by name and margin, the method family
and each graded figure's distance to its bin edge, every semantic fork (SF1 to SF17), mess neutrality per Layer 1
item, the asks' devices and hazards, separation, the distractors' deletion test, the input gates and the realism
sweep. main(target) prints the failures and returns their count."""
import csv
import datetime as dt
import json
import os
import re
import shutil
import sys
import tempfile
import zipfile

import numpy as np
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import params as P  # noqa: E402
import solve as SV  # noqa: E402
import world as W  # noqa: E402

TASK = os.path.dirname(HERE)
CARD = os.path.join(TASK, "..", ".claude", "skills", "fingerprint", "cards", "task118.json")
FLOOR = 1.15                 # every rung's winner by this ratio at least
CALL_FILES = {               # the files the call reads, frozen since the call was built (the asks' re-device kept them)
    "series_new": "bcb7d440cd88affda5a564e747d0875fde88d1e8d5c6b924f038f84e3e2094e4",
    "series_old": "c1adb54c6c2328e9278b9c7a4fea0801f3970fb87ce2b0601099cec2342e8a2e",
    "rates": "12a80a24d1f2ce34154b0bbdece7d263bf7f5386db59f76ff659470f46ed988f",
    "manual": "481fc509292952c6ba4438abff877f1e2a46890632de5c845225f56e21ca62f9",
    "thread": "5a58ea32ab041271e7a79eb90958fd468d2cd1b5c361cbf36ff8e71ef79416a4",
}


def tenth(x):
    return float(np.floor(x * 10.0 + 0.5) / 10.0)


def edge(x):
    """distance of a figure in GBP million to the nearest rounding edge at 0.1"""
    return abs((x * 10.0 - np.floor(x * 10.0)) - 0.5) / 10.0


def board(rr):
    o = sorted(rr, key=lambda c: -rr[c]["gain"])
    return o[0], rr[o[0]]["gain"] / 1e6, o[1], rr[o[1]]["gain"] / 1e6


def appraise_all(S, rt, form="adl10", **kw):
    return {c: SV.appraise(S[c], rt, c, SV.fit(S[c], form), **kw) for c in SV.NAMES}


def texts(target):
    from pypdf import PdfReader
    out = {}
    for k in ("manual", "cpc_note"):
        out[k] = " ".join(p.extract_text() for p in PdfReader(os.path.join(target, P.FILES[k])).pages)
        out[k] = re.sub(r"\s+", " ", out[k])
    out["thread"] = open(os.path.join(target, P.FILES["thread"]), encoding="utf-8").read()
    return out


def headers(target):
    out = {}
    for f in os.listdir(target):
        p = os.path.join(target, f)
        if f.endswith(".csv"):
            with open(p, encoding="utf-8") as fh:
                out[f] = next(csv.reader(fh))
        elif f.endswith(".xlsx"):
            wb = load_workbook(p, read_only=True)
            out[f] = [str(v) for ws in wb.worksheets for v in next(ws.iter_rows(values_only=True)) if v is not None]
    return out


def main(target):
    fails = []
    n = [0]

    def ok(label, cond, detail=""):
        n[0] += 1
        if not cond:
            fails.append("%s %s" % (label, detail))

    R = SV.call(target)
    S, rt, routes = R["S"], R["rt"], R["routes"]
    g = routes["golden"]
    w, wm, r, rm = board(g)
    T = texts(target)

    # ---------------------------------------------------------------- the three rules
    ok("R1 raw material: 75 quarters of clearances and prices per duty from the two layouts",
       all(len(S[c]["q"]) == P.QUARTERS and S[c]["quarters"][0] == (2008, 1) and S[c]["quarters"][-1] == (2026, 3)
           for c in SV.NAMES))
    H = headers(target)
    bad = [(f, h) for f, hs in H.items() for h in hs if re.search(r"elastic|excess|burden|welfare|deadweight", h, re.I)]
    ok("R1 no shipped column holds an elasticity or an excess burden", not bad, bad)
    # R2: every Layer 1 item mishandled leaves the call and the critical components
    def call_of(SS, rr=rt, form="adl10"):
        b = board(appraise_all(SS, rr, form))
        return b[0], tenth(b[1]), b[2], tenth(b[3])
    gold = (w, tenth(wm), r, tenth(rm))
    ok("R2 golden", gold == ("SPIR", 0.9, "CIDR", 0.4), gold)
    newonly = {c: dict(quarters=s["quarters"][P.NEW_FROM:], q=s["q"][P.NEW_FROM:], p=s["p"][P.NEW_FROM:])
               for c, s in S.items()}
    ok("R2 layout: the 2012-2026 release alone (back-series workbook dropped)", call_of(newonly) == gold,
       call_of(newonly))
    noprelim = {c: dict(quarters=s["quarters"][:-2], q=s["q"][:-2], p=s["p"][:-2]) for c, s in S.items()}
    ok("R2 preliminary: 2026 Q2 and Q3 dropped", call_of(noprelim) == gold, call_of(noprelim))
    rt26 = {c: [x for x in v if x[0].year < 2027] for c, v in rt.items()}
    orig = SV.baseline

    def base26(s, rt_, code):
        q0, c, _ = orig(s, rt_, code)
        return q0, c, SV.rate_on(rt_, code, dt.date(2026, 12, 31))
    SV.baseline = base26
    try:
        b26 = call_of(S, rt26)
    finally:
        SV.baseline = orig
    ok("R2 duty rates: the 2027 rates already set missed (2026 rates)", b26 == gold, b26)
    netgst = {c: dict(quarters=s["quarters"], q=s["q"], p=s["p"] / (1 + P.GST)) for c, s in S.items()}
    fits_net = {c: SV.fit(netgst[c]) for c in SV.NAMES}
    bnet = board({c: SV.appraise(S[c], rt, c, fits_net[c]) for c in SV.NAMES})
    ok("R2 GST history: the 3 per cent years taken at 5 per cent (prices stripped at a flat rate before fitting)",
       (bnet[0], tenth(bnet[1]), bnet[2], tenth(bnet[3])) == gold, bnet)
    ok("R2 quarter labels: both layouts parse to the same quarter", SV.qkey("2008 Q1") == SV.qkey("2008Q1") == (2008, 1))
    # R3a / R3b
    b = board(routes["static_triangle"])
    ok("R3a analyst test: static log-log elasticity, the cut's own triangle, every filed rule applied, files cider",
       b[0] == "CIDR" and b[1] / b[3] >= FLOOR, b)
    b = board(routes["static_triangle_static_revenue"])
    ok("R3a the textbook route (revenue before response too) files cider", b[0] == "CIDR" and b[1] / b[3] >= FLOOR, b)
    ok("R3b the decisive step alone (the response within 2027 on the whole wedge) files the golden",
       (w, tenth(wm), r, tenth(rm)) == gold)
    b = board(routes["static_full"])
    ok("R3b the whole wedge alone, static elasticity, files road fuel", b[0] == "FUEL" and b[1] / b[3] >= FLOOR, b)
    b = board(routes["triangle_year"])
    ok("R3b the response within 2027 alone, the cut's own triangle, files cider", b[0] == "CIDR" and b[1] / b[3] >= FLOOR, b)
    blocks = []
    for c in SV.NAMES:
        q = list(np.round(S[c]["q"], 1))
        blocks += [tuple(q[i:i + 4]) for i in range(len(q) - 3)]
    ok("R3c no plant: no repeated four-quarter block in any series", len(blocks) == len(set(blocks)))
    names_hit = [k for k in ("spirits", "cider", "beer", "wine") if k in T["thread"].lower() or k in T["manual"].lower()]
    ok("R3c no plant: neither the manual nor the thread names an alcohol duty", not names_hit, names_hit)

    # ---------------------------------------------------------------- the ladder
    want = {"golden": "SPIR", "family_trend": "SPIR", "family_adl11": "SPIR", "family_linear": "SPIR",
            "static_triangle": "CIDR", "triangle_year": "CIDR", "longrun_full": "FUEL", "static_full": "FUEL",
            "static_triangle_static_revenue": "CIDR", "golden_static_revenue": "SPIR", "longrun_triangle": "FUEL",
            "duty_only": "SPIR", "impact_full": "SPIR", "yearend_full": "SPIR"}
    for k, c in want.items():
        b = board(routes[k])
        ok("ladder %s files %s by %.2f or more" % (k, c, FLOOR), b[0] == c and b[1] / b[3] >= FLOOR, b)
    for k in ("golden_static_revenue", "duty_only", "impact_full", "yearend_full"):
        b = board(routes[k])
        ok("figure rung %s lands spirits at another figure" % k, tenth(b[1]) != 0.9 and edge(b[1]) >= 0.01, b)
    # the method family: every defensible fit and welfare integral files the same four graded figures
    for form in ("adl10", "adl10t", "adl11"):
        for wel in ("full", "linear"):
            b = board(appraise_all(S, rt, form, welfare=wel))
            ok("family %s %s converges" % (form, wel), (b[0], tenth(b[1]), b[2], tenth(b[3])) == gold, b)
            ok("family %s %s spirits figure 0.015 or more from a bin edge" % (form, wel), edge(b[1]) >= 0.015, b[1])
            ok("family %s %s cider figure 0.015 or more from a bin edge" % (form, wel), edge(b[3]) >= 0.015, b[3])

    # ---------------------------------------------------------------- semantic forks
    for c in SV.NAMES:
        x = g[c]
        p1 = x["p1"]
        rev = float(np.sum((x["p0"] - x["c"]) * x["q0"] - (p1 - x["c"]) * x["q1"]))
        ok("SF3 the envelope: GBP 1.5 million of duty and GST given up after the response, %s" % c,
           abs(rev - P.ENVELOPE) < 1.0, rev)
        ok("SF5 full pass-through with GST on top, %s" % c, abs(p1 - (x["c"] + x["d27"] - x["cut"]) * (1 + P.GST)) < 1e-9)
        ok("SF6 the 2027 rate already set, %s" % c, x["d27"] == SV.rate_on(rt, c, dt.date(2027, 1, 1)))
        d_now = SV.rate_on(rt, c, dt.date(2026, 9, 1))
        ok("SF7 producer price: the 2026 Q3 shelf price net of GST less the duty then in force, %s" % c,
           abs(x["c"] - (S[c]["p"][-1] / (1 + P.GST) - d_now)) < 1e-9)
        ok("SF8 the wedge is duty and GST, %s" % c, abs((x["p0"] - x["c"]) - (x["d27"] * (1 + P.GST) + P.GST * x["c"])) < 1e-9)
        ok("SF1 the Budget year: four quarters of 2027, %s" % c, len(x["eps"]) == 4 and len(x["q0"]) == 4)
    ok("SF4 the starting point: 2025 Q4 to 2026 Q3, preliminary quarters included",
       all(S[c]["quarters"][-4:] == [(2025, 4), (2026, 1), (2026, 2), (2026, 3)] for c in SV.NAMES))
    ok("SF2 the next best is the second duty under the same appraisal", R["runner"] == "CIDR" and R["winner"] == "SPIR")

    # ---------------------------------------------------------------- the asks
    modes = {m: SV.asks(target, m) for m in SV.MODES}
    A = modes["golden"]
    Sw = W.series()
    _, truth, whs = W.declarations(Sw, W.rate_history(Sw))
    ok("asks: the golden reproduces the generator's truth to the penny",
       all(abs(A[c][q] - truth[c][q]) < 0.006 for c in truth for q in truth[c]))
    cells = [(c, q) for c in SV.NAMES for q in P.ASK_QUARTERS]

    def off(m):
        return sum(SV.thousand(modes[m][c][q]) != SV.thousand(A[c][q]) for c, q in cells)

    def shift(m):
        return min(abs(modes[m][c][q] - A[c][q]) for c, q in cells)
    # the primary: corrections lodged as new declarations naming the one they correct, chained across quarter ends
    for m in ("one_hop", "one_hop_duty", "own_quarter", "latest_per_named", "every_declaration", "named_by_any"):
        ok("SF10 the careful reading %s keeps at most 8 of 20 cells, each miss GBP 500 or more off" % m,
           20 - off(m) <= 8 and shift(m) >= 500, (20 - off(m), round(shift(m))))
    ok("SF10 the chain read right but each correction priced on its own day moves the year-end cells only",
       0 < off("first_own_rate") <= 6, off("first_own_rate"))
    ok("SF16 the repeated declaration duty kept once, on its first item, misses every cell",
       off("header_once") == 20, off("header_once"))
    ok("SF16 the repeated duty shared by litres misses every cell", off("header_litres") == 20, off("header_litres"))
    ok("SF16 the repeated duty summed on every item row misses every cell", off("header_every_item") == 20)
    ok("SF9 Eskerholm date: UTC dates move 10 or more cells", off("utc_day") >= 10, off("utc_day"))
    ok("SF12 repayments: leaving them out moves 16 or more cells", off("no_repayments") >= 16, off("no_repayments"))
    ok("SF13 removals to lapsed premises: leaving them out moves 16 or more cells", off("no_lapsed") >= 16,
       off("no_lapsed"))
    ok("SF14 releases are procedure 4071: adding the 4000 imports misses every cell", off("with_imports") == 20)
    ok("asks: exact duplicate rows kept moves 16 or more cells", off("no_dedup") >= 16, off("no_dedup"))
    rows = SV.read(os.path.join(target, P.FILES["declarations"]))
    canc = [x for x in rows if x["status"] == "CANCELLED"]
    ok("SF11 cancellations ship and carry duty", len(canc) > 100 and all(float(x["duty_gbp"]) > 0 for x in canc
                                                                        if x["cpc"] == "4071"))
    ok("SF15 marked gas oil (REBF) rides on road fuel releases and is not one of the five",
       any(x["duty_code"] == "REBF" for x in rows) and set(A) == set(SV.NAMES))
    # the chains, read from the shipped rows: every correction names a release declaration in the extract, none is
    # corrected twice, no chain starts at a cancelled release, and corrections of corrections cross every quarter end
    seen, decl = set(), {}
    for x in rows:
        k = tuple(x.values())
        if k in seen:
            continue
        seen.add(k)
        if x["cpc"] == "4071":
            decl.setdefault(x["declaration_no"], []).append(x)
    named = {k: v[0]["related_declaration"] for k, v in decl.items() if v[0]["related_declaration"]}
    ok("SF10 every correction names a release declaration in the extract", all(nm in decl for nm in named.values()))
    ok("SF10 no declaration is corrected twice", len(set(named.values())) == len(named))
    ok("SF10 no corrected declaration is cancelled", all(decl[nm][0]["status"] != "CANCELLED" for nm in named.values()))
    rep = [x for x in rows if x["cpc"] == "6171"]
    ok("SF12 every repayment names a release that no correction names, its first declaration",
       all(x["related_declaration"] in decl and x["related_declaration"] not in named for x in rep))

    def first(k):
        while k in named:
            k = named[k]
        return k
    crossing = {}
    for k, nm in named.items():
        if nm in named:
            f = first(k)
            qf = SV.quarter(SV.local_day(decl[f][0]["accepted_at_utc"]))
            qn = SV.quarter(SV.local_day(decl[nm][0]["accepted_at_utc"]))
            if qn != qf:
                crossing[(decl[k][0]["duty_code"], qn)] = crossing.get((decl[k][0]["duty_code"], qn), 0) + 1
    ok("SF10 corrections of corrections whose named declaration sits a quarter after the release, for every duty in "
       "every quarter of the window", all(crossing.get(cq, 0) >= 3 for cq in cells),
       [cq for cq in cells if crossing.get(cq, 0) < 3])
    # the forcing figure: every release declaration's duty equals its items at the rates of the day the release was
    # first accepted, on every item row
    bad, multi = 0, 0
    for k, items in decl.items():
        day0 = SV.local_day(decl[first(k)][0]["accepted_at_utc"])
        multi += len(items) > 1
        rate = lambda c: P.REBATED_RATE if c == "REBF" else SV.rate_on(rt, c, day0)
        tot = sum(round(float(i["quantity_litres"]) * rate(i["duty_code"]), 2) for i in items)
        bad += abs(tot - float(items[0]["duty_gbp"])) > 0.005 or len({i["duty_gbp"] for i in items}) != 1
    ok("SF16 every release declaration's duty_gbp is its items at the rates of the release's first day, on every row",
       bad == 0 and multi > 1000, (bad, multi))
    lapsed = {wx["id"]: wx for wx in whs if wx["id"] in W.LAPSED}
    on_end = [x for x in rows if x["cpc"] == "7171" and x["warehouse_to"] in lapsed
              and SV.local_day(x["accepted_at_utc"]) == lapsed[x["warehouse_to"]]["ended"]]
    ok("SF13 no removal to a lapsed premises on its end date", not on_end)
    late = {}
    for x in rows:
        if x["cpc"] == "7171" and x["warehouse_to"] in lapsed:
            d = SV.local_day(x["accepted_at_utc"])
            if d > lapsed[x["warehouse_to"]]["ended"]:
                late.setdefault((x["duty_code"], SV.quarter(d)), 0)
                late[(x["duty_code"], SV.quarter(d))] += 1
    ok("SF13 every duty and quarter carries removals to a lapsed premises", all(late.get(cq, 0) >= 2 for cq in cells),
       [cq for cq in cells if late.get(cq, 0) < 2])
    margins = [abs((A[c][q] % 100.0) - 50.0) for c, q in cells]
    ok("SF17 every excise cell GBP 1 or more from a rounding edge at GBP 100", min(margins) >= 1.0, min(margins))
    ok("SF17 the excess burden figures are graded to GBP 0.1 million", edge(wm) >= 0.015 and edge(rm) >= 0.015)
    for form in ("adl10", "adl10t", "adl11"):
        fs = {c: SV.fit(S[c], form) for c in ("SPIR", "FUEL")}
        made = {c: SV.path(fs[c])[3] / SV.path(fs[c], "longrun")[0] for c in fs}
        ok("prose (%s): by 2027's fourth quarter spirits make nearly all their long-run response, road fuel under a "
           "third" % form, made["SPIR"] >= 0.95 and made["FUEL"] < 0.33, made)
        ga = appraise_all(S, rt, form)
        rel = {c: ga[c]["gain"] / ga["SPIR"]["gain"] for c in ("BEER", "WINE", "FUEL")}
        ok("prose (%s): beer and wine each about a quarter of spirits, road fuel under a fifth" % form,
           0.2 <= rel["BEER"] <= 0.3 and 0.2 <= rel["WINE"] <= 0.3 and rel["FUEL"] < 0.2, rel)
    ok("prose: spirits carry the largest wedge of the alcohol duties, road fuel a larger one",
       max(("BEER", "CIDR", "WINE", "SPIR"), key=lambda c: g[c]["wedge_share"]) == "SPIR" and
       g["FUEL"]["wedge_share"] > g["SPIR"]["wedge_share"])
    ok("prose: cider is the most price-responsive duty in every quarter of 2027",
       all(g["CIDR"]["eps"][k] == min(g[c]["eps"][k] for c in SV.NAMES) for k in range(4)))
    shares = [appraise_all(S, rt, f)["SPIR"]["cut"] / g["SPIR"]["d27"] for f in ("adl10", "adl10t", "adl11")]
    ok("SF17 the spirits cut is about a fifth of its 2027 rate across the family", all(0.19 <= x <= 0.22 for x in shares),
       shares)
    ok("asks: the CE14 text files UK time, the original's period, repayments, removals and the end day",
       all(s in T["cpc_note"] for s in ("United Kingdom time", "in the period in which the release was first accepted",
                                        "credited in the period the repayment is accepted",
                                        "whatever procedure the removal was declared under",
                                        "up to the day before its end date")))

    import hashlib
    for k, digest in CALL_FILES.items():
        got = hashlib.sha256(open(os.path.join(target, P.FILES[k]), "rb").read()).hexdigest()
        ok("freeze: %s is byte-identical to the call's build" % P.FILES[k], got == digest, got[:12])

    # ---------------------------------------------------------------- separation
    tmp = tempfile.mkdtemp()
    try:
        t2 = os.path.join(tmp, "t")
        shutil.copytree(target, t2)
        for k in ("declarations", "cpc", "cpc_note", "warehouses"):
            os.remove(os.path.join(t2, P.FILES[k]))
        R2 = SV.call(t2)
        ok("separation: the call is unchanged with every ask file deleted",
           board(R2["routes"]["golden"]) == board(g))
        for k in P.DISTRACTORS:
            t3 = os.path.join(tmp, "d_" + k)
            shutil.copytree(target, t3)
            os.remove(os.path.join(t3, P.FILES[k]))
            ok("distractor %s deleted: the call unchanged" % k, board(SV.call(t3)["routes"]["golden"]) == board(g))
            ok("distractor %s deleted: every ask unchanged" % k, SV.asks(t3) == A)
    finally:
        shutil.rmtree(tmp)
    # the published clearances and Customs' duty-paid litres agree loosely and never tie
    lit = {}
    for k, items in decl.items():
        if k in set(named.values()) or decl[first(k)][0]["status"] == "CANCELLED":
            continue
        q = SV.quarter(SV.local_day(decl[first(k)][0]["accepted_at_utc"]))
        for i in items:
            lit[(i["duty_code"], q)] = lit.get((i["duty_code"], q), 0.0) + float(i["quantity_litres"])
    seen = set()
    for x in rows:
        k = tuple(x.values())
        if k in seen or x["cpc"] != "4000":
            continue
        seen.add(k)
        q = SV.quarter(SV.local_day(x["accepted_at_utc"]))
        lit[(x["duty_code"], q)] = lit.get((x["duty_code"], q), 0.0) + float(x["quantity_litres"])
    dev = []
    for c, q in cells:
        i = S[c]["quarters"].index((int(q[:4]), int(q[-1])))
        dev.append(lit[(c, q)] / (S[c]["q"][i] * 1000.0) - 1.0)
    ok("separation: Customs' duty-paid litres within 8 per cent of the clearances and never tying",
       max(abs(x) for x in dev) < 0.08 and min(abs(x) for x in dev) > 0.0005, (min(dev), max(dev)))

    # ---------------------------------------------------------------- input gates and realism
    files = sorted(os.listdir(target))
    ok("gate: 10 or more files", len(files) >= 10, len(files))
    ok("gate: 3 or more formats", len({f.rsplit(".", 1)[1] for f in files}) >= 3)
    big = max(sum(1 for _ in open(os.path.join(target, f), encoding="utf-8")) - 1 for f in files if f.endswith(".csv"))
    ok("gate: one file of 25,000 rows or more", big >= 25000, big)
    meta = json.load(open(os.path.join(TASK, "metadata.json")))
    ok("gate: two or more distractors named in metadata", len(meta["distractor_files"]) >= 2 and
       all(f in files for f in meta["distractor_files"]))
    ok("gate: no file name says distractor", not any("distract" in f for f in files))
    card = json.load(open(CARD))
    ok("realism: personas are the card's", sorted(P.PEOPLE.values()) == sorted(card["world"]["people"]))
    rep = [(c, i) for c in SV.NAMES for i in range(1, P.QUARTERS)
           if round(float(S[c]["q"][i]), 1) == round(float(S[c]["q"][i - 1]), 1)]
    ok("realism: no clearance repeated in consecutive quarters", not rep, rep)
    same = [i for i in range(P.QUARTERS) if len({round(float(S[c]["q"][i]), 1) for c in SV.NAMES}) < len(SV.NAMES)]
    ok("realism: no two duties publish the same clearance in a quarter", not same, same)
    ok("realism: no round excise cell", all(A[c][q] % 1000 != 0 for c, q in cells))
    acc = [SV.local_day(x["accepted_at_utc"]) for x in rows]
    ok("realism: every acceptance inside the extract's local dates", min(acc) >= P.DECL_FROM and max(acc) <= P.DECL_TO)
    hours = {}
    for x in rows:
        t = dt.datetime.strptime(x["accepted_at_utc"], "%Y-%m-%dT%H:%M:%SZ")
        hours[t.hour] = hours.get(t.hour, 0) + 1
    ok("realism: acceptance hours spread over the working day and the overnight run", len(hours) >= 14)
    ok("realism: the thread is dated before the as-of date", "19 Oct 2026" in T["thread"])
    for f in files:
        if f.endswith((".xlsx", ".docx")):
            z = zipfile.ZipFile(os.path.join(target, f))
            core = z.read("docProps/core.xml").decode()
            ok("realism: %s carries no writing library's name" % f, "openpyxl" not in core.lower() and
               "python" not in core.lower())

    print("checks: %d assertions, %d failed" % (n[0], len(fails)))
    for f in fails:
        print("FAIL", f)
    return len(fails)


if __name__ == "__main__":
    sys.exit(1 if main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(TASK, "target")) else 0)
