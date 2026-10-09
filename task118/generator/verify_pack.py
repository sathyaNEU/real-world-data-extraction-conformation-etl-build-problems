"""Independent verifier for task118: python3 verify_pack.py <target_dir>

Reads only the shipped files in target_dir. It imports nothing from the generator, reads no seed, parameter
or side file, and recomputes every rung, every rival the change log refutes, every calibration outcome and
every graded figure on its own code path (DuckDB over the parquet, a profile-likelihood prior fit, a walk of
each document's published states in the CMS revisions that skips a note the document had already shown and pairs
a note on an unchanged headline with the bare headline change in the 20 minutes before it, else in the 20 minutes
after it, else with the entry fix in the 20 minutes before it, each panel release read on the section list it was
issued on and counted only for the months the release log lists for it, each section's month taken from the latest
such release). The CLAIMS block is the answer key it checks.
"""
import csv
import datetime as dt
import math
import re
import statistics
import sys
import zipfile
from collections import defaultdict
from pathlib import Path

import duckdb
import numpy as np
from openpyxl import load_workbook
from scipy.optimize import minimize_scalar

CLAIMS = {
    "call": "CUL-N", "call_figure": 1_250_000, "runner_up": "LOC-M", "gap": 450_000,
    "rung_leaders": ["POL-N", "SPT-M", "BUS-N", "LOC-M", "CUL-N"],
    "extra": {"POL-N": 550_000, "SPT-M": 100_000, "BUS-N": 700_000, "SPT-N": 650_000, "LOC-M": 800_000,
              "CUL-N": 1_250_000},
    "readers": {"POL-N": 3_180_000, "SPT-M": 470_000, "BUS-N": 1_760_000, "SPT-N": 2_931_000, "LOC-M": 617_000,
                "CUL-N": 1_412_000},
    "per_reader": {"POL-N": 0.2, "SPT-M": 0.2, "BUS-N": 0.4, "SPT-N": 0.2, "LOC-M": 1.3, "CUL-N": 0.9},
    "corrections": {"POL-N": 47, "SPT-M": 18, "BUS-N": 39, "SPT-N": 42, "LOC-M": 34, "CUL-N": 14},
    "rate": {"POL-N": 5.2, "SPT-M": 7.3, "BUS-N": 4.9, "SPT-N": 3.8, "LOC-M": 4.9, "CUL-N": 3.6},
    "median_minutes": {"POL-N": 86, "SPT-M": 55, "BUS-N": 89, "SPT-N": 47, "LOC-M": 102, "CUL-N": 119},
    "backtest_hits": 7,
    "bulletin_march": (11, 22),
}
SHORTLIST = ["POL-N", "SPT-M", "BUS-N", "SPT-N", "LOC-M", "CUL-N"]
PLATFORM = ("home_web", "section_web", "feed_app", "section_app", "related_links")

results = []


def check(name, ok, detail=""):
    results.append((name, bool(ok), detail))


def r50k(x):
    return int(math.floor(x / 50_000 + 0.5) * 50_000)


def one_dp(x):
    return math.floor(x * 10 + 0.5) / 10


def main(target):
    t = Path(target)
    files = {p.name: p for p in t.iterdir()}

    def find(prefix):
        hit = [n for n in files if n.startswith(prefix)]
        assert len(hit) == 1, prefix
        return files[hit[0]]

    spine = find("pageviews_by_source_age")
    archive = find("headline_tests_archive")
    staff = find("newsroom_staff_list")
    plan = find("audience_plan_2027")
    changelog = find("headline_squad_change_log")
    desks_f = find("desk_register")
    dash_f = find("experimentation_dashboard_export")
    cms_f = find("cms_revisions_web_desks")
    panel_f = find("panel_monthly_audience")
    panelwb = find("panel_reference_workbook")

    # ------------------------------------------------------------------ desks, plan, staff
    desks = {r["desk_code"]: r for r in csv.DictReader(open(desks_f, encoding="utf-8"))}
    wb = load_workbook(plan, read_only=True, data_only=True)
    rows = list(wb["Plan 2027"].iter_rows(values_only=True))
    planned = {r[0]: int(r[4]) for r in rows[1:] if r[0] and r[0] != "Total"}
    wb = load_workbook(staff, read_only=True, data_only=True)
    srows = list(wb["Staff"].iter_rows(values_only=True))
    team = {r[0]: r[2] for r in srows[1:]}

    # ------------------------------------------------------------------ the archive, test by test
    tests = defaultdict(list)
    meta = {}
    for r in csv.DictReader(open(archive, encoding="utf-8")):
        tests[r["test_id"]].append(r)
        meta[r["test_id"]] = (r["engine"], r["desk_code"], int(r["article_id"]), r["owner_staff_id"],
                              r["started_at"], r["concluded_at"])
    L, S2, pkg_test = [], [], []
    per_test = {}
    for tid, pk in tests.items():
        ctl = [p for p in pk if p["variant"] == "control"][0]
        n0, c0 = int(ctl["impressions"]), int(ctl["clicks"])
        p0 = c0 / n0
        lifts = []
        for p in pk:
            if p["variant"] == "control":
                continue
            n, c = int(p["impressions"]), int(p["clicks"])
            pv = c / n
            rr = pv / p0
            s2 = rr * rr * ((1 - pv) / (n * pv) + (1 - p0) / (n0 * p0))
            L.append(rr - 1)
            S2.append(s2)
            pkg_test.append(tid)
            lifts.append((rr - 1, s2, p["shipped"] == "Y", pv - p0))
        per_test[tid] = lifts
    L, S2 = np.array(L), np.array(S2)

    def neg_profile(tau2):
        w = 1 / (tau2 + S2)
        mu = np.sum(w * L) / np.sum(w)
        return 0.5 * np.sum(np.log(tau2 + S2) + (L - mu) ** 2 * w)

    res = minimize_scalar(neg_profile, bounds=(1e-7, 0.05), method="bounded", options={"xatol": 1e-13})
    tau2 = float(res.x)
    w = 1 / (tau2 + S2)
    mu = float(np.sum(w * L) / np.sum(w))
    prior = (mu, tau2)
    mom = (float(L.mean()), float(L.var() - S2.mean()))

    def win(tid, pr, shrink=True):
        for l, s2, shipped, _ in per_test[tid]:
            if shipped:
                return (pr[0] + pr[1] / (pr[1] + s2) * (l - pr[0])) if shrink else l
        return 0.0

    by_desk = defaultdict(list)
    for tid, m in meta.items():
        by_desk[m[1]].append(tid)
    web = {d: [x for x in by_desk[d] if meta[x][0] == "web"] for d in by_desk}
    raw = {d: np.mean([win(x, prior, False) for x in web[d]]) for d in SHORTLIST}
    shr = {d: np.mean([win(x, prior) for x in web[d]]) for d in SHORTLIST}
    check("one test per article", len({m[2] for m in meta.values()}) == len(meta))
    check("every shortlisted-desk test is owned by that desk's own staff",
          all(team.get(meta[x][3]) == meta[x][1] for d in SHORTLIST for x in by_desk[d]))
    check("every app-desk test is owned by the headline squad",
          all(team.get(m[3]) == "AUD-HS" for m in meta.values() if desks[m[1]]["distribution"] == "app only"))

    # ------------------------------------------------------------------ the spine, through DuckDB
    con = duckdb.connect()
    con.execute("CREATE TABLE pv AS SELECT * FROM read_parquet('%s')" % str(spine).replace("'", "''"))
    tested = {m[2] for m in meta.values()}
    won = {meta[x][2] for x in meta if any(sh for _, _, sh, _ in per_test[x])}
    con.execute("CREATE TABLE tested AS SELECT unnest(?::BIGINT[]) AS article_id", [sorted(tested)])
    con.execute("CREATE TABLE won AS SELECT unnest(?::BIGINT[]) AS article_id", [sorted(won)])
    plat_list = ",".join("'%s'" % s for s in PLATFORM)
    q = """
      SELECT p.desk_code,
             sum(pageviews) AS total,
             sum(pageviews) FILTER (WHERE source_code IN ({pl})) AS plat,
             sum(pageviews) FILTER (WHERE source_code = 'partner_apps') AS part,
             sum(pageviews) FILTER (WHERE t.article_id IS NOT NULL) AS t_all,
             sum(pageviews) FILTER (WHERE t.article_id IS NOT NULL AND source_code IN ({pl})) AS t_plat,
             sum(pageviews) FILTER (WHERE t.article_id IS NOT NULL AND source_code <> 'partner_apps') AS t_np,
             sum(pageviews) FILTER (WHERE w.article_id IS NOT NULL) AS w_all,
             sum(pageviews) FILTER (WHERE w.article_id IS NOT NULL AND source_code IN ({pl})) AS w_plat,
             sum(pageviews) FILTER (WHERE w.article_id IS NOT NULL AND source_code <> 'partner_apps') AS w_np,
             count(DISTINCT p.article_id) AS n_art,
             count(DISTINCT p.article_id) FILTER (WHERE t.article_id IS NOT NULL) AS n_tested,
             count(DISTINCT p.article_id) FILTER (WHERE published_date >= DATE '2025-10-01') AS n_win
      FROM pv p LEFT JOIN tested t USING (article_id) LEFT JOIN won w USING (article_id)
      GROUP BY p.desk_code""".format(pl=plat_list)
    C = {r[0]: dict(zip(["total", "plat", "part", "t_all", "t_plat", "t_np", "w_all", "w_plat", "w_np", "n_art",
                         "n_tested", "n_win"], [x or 0 for x in r[1:]])) for r in con.execute(q).fetchall()}
    check("spine totals equal the plan's 2027 clicks for every desk", all(C[d]["total"] == planned[d] for d in planned))
    dash = {}
    lines = open(dash_f, encoding="utf-8").read().splitlines()
    hdr = lines.index("vertical,tests_concluded,tests_with_variant_shipped,avg_winning_lift_pct")
    for r in csv.reader(lines[hdr + 1:]):
        dash[r[0]] = float(r[3]) / 100
    # the dashboard against the archive
    lo, hi = dt.datetime(2025, 9, 30, 14), dt.datetime(2026, 9, 30, 14)
    okd = True
    for v, lift in dash.items():
        ids = [x for x, m in meta.items() if desks[m[1]]["vertical"] == v
               and lo <= dt.datetime.strptime(m[5], "%Y-%m-%dT%H:%M:%SZ") < hi]
        okd &= round(100 * np.mean([win(x, prior, False) for x in ids]), 2) == round(100 * lift, 2)
    check("dashboard reproduces from the archive", okd)

    vert = {d: desks[d]["vertical"] for d in desks}
    nweb = {d: len(web.get(d, [])) for d in desks}
    shr_all = {d: np.mean([win(x, prior) for x in web[d]]) for d in desks if web.get(d)}
    pooled = {d: sum(nweb[x] * shr_all[x] for x in shr_all if vert[x] == vert[d]) /
              sum(nweb[x] for x in shr_all if vert[x] == vert[d]) for d in SHORTLIST}
    lifts = {"dash": {d: dash[vert[d]] for d in SHORTLIST}, "raw": raw, "shr": shr, "pool": pooled}

    def cellv(lb, d, base, net):
        c = C[d]
        b = {"all": c["total"], "np": c["total"] - c["part"], "plat": c["plat"]}[base]
        f = {"none": 1.0,
             "click": 1 - {"all": c["t_all"], "np": c["t_np"], "plat": c["t_plat"]}[base] / b,
             "count": 1 - c["n_tested"] / c["n_art"],
             "won": 1 - {"all": c["w_all"], "np": c["w_np"], "plat": c["w_plat"]}[base] / b}[net]
        return lifts[lb][d] * b * f

    rungs = [("dash", "all", "none"), ("raw", "all", "none"), ("shr", "all", "none"), ("shr", "plat", "none"),
             ("shr", "plat", "click")]
    rv = []
    for k, (lb, b, n) in enumerate(rungs):
        vals = {d: cellv(lb, d, b, n) for d in SHORTLIST}
        o = sorted(vals, key=vals.get, reverse=True)
        rv.append(vals)
        check("rung %d leader %s, margin at least 1.2" % (k, CLAIMS["rung_leaders"][k]),
              o[0] == CLAIMS["rung_leaders"][k] and vals[o[0]] / vals[o[1]] >= 1.2,
              "%s %.3f" % (o[0], vals[o[0]] / vals[o[1]]))
    r4 = rv[4]
    o4 = sorted(r4, key=r4.get, reverse=True)
    check("call: %s at %s" % (CLAIMS["call"], CLAIMS["call_figure"]), o4[0] == CLAIMS["call"]
          and r50k(r4[o4[0]]) == CLAIMS["call_figure"], "%.0f" % r4[o4[0]])
    check("runner-up and gap", o4[1] == CLAIMS["runner_up"] and r50k(r4[o4[0]] - r4[o4[1]]) == CLAIMS["gap"],
          "%s %.0f" % (o4[1], r4[o4[0]] - r4[o4[1]]))
    for d in SHORTLIST:
        check("extra clicks %s" % d, r50k(r4[d]) == CLAIMS["extra"][d], "%.0f" % r4[d])
    grid_cul = []
    for lb in lifts:
        for b in ("all", "np", "plat"):
            for n in ("none", "click", "count", "won"):
                vals = {d: cellv(lb, d, b, n) for d in SHORTLIST}
                if max(vals, key=vals.get) == "CUL-N":
                    grid_cul.append((lb, b, n))
    check("only the answer cell and its vertical-pooled twin name Culture across the 48-cell grid",
          sorted(grid_cul) == [("pool", "plat", "click"), ("shr", "plat", "click")], str(grid_cul))
    unt = {d: 1 - C[d]["t_plat"] / C[d]["plat"] for d in SHORTLIST}
    check("dominance over Local-metro at least 1.2",
          (unt["CUL-N"] / unt["LOC-M"]) / (rv[3]["LOC-M"] / rv[3]["CUL-N"]) >= 1.2)

    # ------------------------------------------------------------------ the change log back-test
    wb = load_workbook(changelog, read_only=True, data_only=True)
    crow = list(wb["Embeddings"].iter_rows(values_only=True))
    hdrc = crow[0]
    emb = [dict(zip(hdrc, r)) for r in crow[1:] if r[0] is not None]
    app_tests = defaultdict(list)
    for x, m in meta.items():
        if m[0] == "app":
            app_tests[(m[1], int(m[4][:4]))].append(x)
    imp = {x: sum(int(p["impressions"]) for p in tests[x]) for x in tests}
    clk = {x: sum(int(p["clicks"]) for p in tests[x]) for x in tests}
    first2h = dict(con.execute("SELECT desk_code, sum(pageviews) FILTER (WHERE age_band='0-2h') / sum(pageviews) "
                               "FROM pv GROUP BY desk_code").fetchall())

    def fit_mom(ids):
        sel = np.isin(np.array(pkg_test), ids)
        return float(L[sel].mean()), max(float(L[sel].var() - S2[sel].mean()), 0.0)

    preds = defaultdict(dict)
    real = {}
    for e in emb:
        ids = app_tests[(e["desk_code"], e["started"].year)]
        real[e["embedding"]] = float(e["realised_incremental_clicks"])
        plan_c = float(e["planned_clicks_m"]) * 1e6
        check("change log row %d: tests and observed lift reproduce" % e["embedding"],
              len(ids) == e["tests_run"] and round(100 * np.mean([win(x, prior, False) for x in ids]), 1) == round(e["avg_winning_lift_pct"], 1))
        g = np.mean([win(x, prior) for x in ids])
        preds["golden"][e["embedding"]] = g * plan_c
        preds["raw"][e["embedding"]] = np.mean([win(x, prior, False) for x in ids]) * plan_c
        preds["pooled MoM"][e["embedding"]] = np.mean([win(x, mom) for x in ids]) * plan_c
        desk_ids = [x for (dk, y), v in app_tests.items() if dk == e["desk_code"] for x in v]
        preds["per-desk MoM"][e["embedding"]] = np.mean([win(x, fit_mom(desk_ids)) for x in ids]) * plan_c
        preds["per-embedding MoM"][e["embedding"]] = np.mean([win(x, fit_mom(ids)) for x in ids]) * plan_c
        vids = [x for (dk, y), v in app_tests.items() if vert[dk] == vert[e["desk_code"]] for x in v]
        preds["per-vertical MoM"][e["embedding"]] = np.mean([win(x, fit_mom(vids)) for x in ids]) * plan_c
        sv = []
        for x in ids:
            best = max(prior[0] + prior[1] / (prior[1] + s2) * (l - prior[0]) for l, s2, _, _ in per_test[x])
            sv.append(max(0.0, best))
        preds["shrink then max"][e["embedding"]] = np.mean(sv) * plan_c
        wv = np.array([win(x, prior) for x in ids])
        preds["impression-weighted"][e["embedding"]] = np.average(wv, weights=[imp[x] for x in ids]) * plan_c
        preds["click-weighted"][e["embedding"]] = np.average(wv, weights=[clk[x] for x in ids]) * plan_c
        preds["post-test"][e["embedding"]] = g * plan_c * (1 - first2h[e["desk_code"]])
        preds["absolute points"][e["embedding"]] = np.mean(
            [next((dp for _, _, sh, dp in per_test[x] if sh), 0.0) for x in ids]) * plan_c
    rawv = np.array([preds["raw"][k] for k in sorted(real)])
    realv = np.array([real[k] for k in sorted(real)])
    kcut = float(np.sum(rawv * realv) / np.sum(rawv ** 2))
    preds["global haircut"] = {k: kcut * preds["raw"][k] for k in real}
    for name, p in preds.items():
        err = [abs(p[k] / real[k] - 1) for k in real]
        hits = sum(e <= 0.02 for e in err)
        miss = sum(e > 0.03 for e in err)
        if name == "golden":
            check("back-test: the golden method gets %d of 7 within 2%%" % CLAIMS["backtest_hits"],
                  hits == CLAIMS["backtest_hits"], "worst %.2f%%" % (100 * max(err)))
        else:
            check("back-test refutes %s (2 or more of 7 missed by over 3%%)" % name, miss >= 2,
                  "%d hits %d misses" % (hits, miss))
    t3 = [e for e in emb if e["embedding"] == 3][0]
    t6 = [e for e in emb if e["embedding"] == 6][0]
    check("twins identical on every change-log column a lookup sees, realised about 2x apart",
          all(t3[k] == t6[k] for k in ("tests_run", "avg_winning_lift_pct", "planned_clicks_m", "distribution"))
          and 1.95 <= real[3] / real[6] <= 2.1, "%.3f" % (real[3] / real[6]))

    # ------------------------------------------------------------------ readers: each release on its own section list
    wb = load_workbook(panelwb, read_only=True, data_only=True)
    hist = list(wb["Section history"].iter_rows(values_only=True))[1:]
    rlog = [r for r in list(wb["Release log"].iter_rows(values_only=True))[1:] if r[0]]
    rel = {r[0]: r[1] for r in rlog}
    rel_months = {r[0]: {x.strip() for x in str(r[2]).split(",")} for r in rlog}
    first_2026 = [r[1] for r in rlog if (r[3] or "").startswith("First release on the 2026 content taxonomy")][0]
    later_list = {(h[0], h[1]): (h[2], h[5]) for h in hist if h[4] is None}
    earlier_list = {(h[0], h[1]): (h[2], h[5]) for h in hist if h[4] is not None}
    best = {}
    off_log = 0
    for r in csv.DictReader(open(panel_f, encoding="utf-8")):
        if r["period"] not in rel_months[r["release"]]:
            off_log += 1        # a row a release's file carried for a month that release did not publish
            continue
        lst = later_list if rel[r["release"]] >= first_2026 else earlier_list
        name, desk = lst[(r["site_code"], int(r["section_code"]))]
        key = (r["period"], r["site_code"], name)
        if key not in best or rel[r["release"]] > rel[best[key][0]["release"]]:
            best[key] = (r, desk)
    monthly = defaultdict(dict)
    for (period, site, name), (r, desk) in best.items():
        if desk:
            monthly[desk][period] = int(r["unique_audience"])
    readers = {d: statistics.mean(monthly[d].values()) for d in SHORTLIST}
    check("panel rows outside their release's published months are set aside (the July file's April to June rows)",
          off_log == 27, str(off_log))
    for d in SHORTLIST:
        check("readers %s (12 months)" % d, len(monthly[d]) == 12 and round(readers[d], -3) == CLAIMS["readers"][d],
              "%.1f" % readers[d])
        check("extra clicks per reader %s" % d, one_dp(r4[d] / readers[d]) == CLAIMS["per_reader"][d],
              "%.4f" % (r4[d] / readers[d]))

    # ------------------------------------------------------------------ corrections: a walk of the published states
    docs = defaultdict(list)
    info = {}
    for r in csv.DictReader(open(cms_f, encoding="utf-8")):
        docs[r["doc_id"]].append((int(r["revision"]), r["saved_at"], r["status"], r["publish_at"], r["headline_sha1"],
                                  r["correction_note"]))
        info[r["doc_id"]] = (r["desk_code"], r["doc_type"], r["parent_doc"], r["restored_from_doc"], r["migrated_from"])
    copy_of = {v[3]: k for k, v in info.items() if v[3]}
    ts = lambda x: dt.datetime.strptime(x, "%Y-%m-%dT%H:%MZ")        # noqa: E731
    # an entry that publishes a changed headline, by the live blog it belongs to
    entry_fix = defaultdict(set)
    for did, (desk, typ, parent, restored, migrated) in info.items():
        if typ != "post" or restored:
            continue
        revs = sorted(docs[did]) + (sorted(docs[copy_of[did]]) if did in copy_of else [])
        shown = [r for r in revs if r[2] == "live"]
        blog = info[parent][3] or parent
        for a, b in zip(shown, shown[1:]):
            if b[4] != a[4]:
                entry_fix[blog].add(ts(b[1]))
    count, arts, mins, march = defaultdict(int), defaultdict(int), defaultdict(list), [0, 0]
    m_lo, m_hi = dt.datetime(2026, 2, 28, 14), dt.datetime(2026, 3, 31, 14)      # March 2026, AEST
    for did, revs in docs.items():
        desk, typ, parent, restored, migrated = info[did]
        if restored or migrated or typ == "post":
            continue
        chain = sorted(revs)
        if did in copy_of:
            chain = chain + sorted(docs[copy_of[did]])
        live = [r for r in chain if r[2] == "live"]
        if not live:
            continue
        arts[desk] += 1
        states, golive = live, ts(live[0][1])
        sched = [r for r in chain if r[2] == "scheduled" and r[0] < live[0][0]]
        if sched and sched[-1][3] and ts(sched[-1][3]) < golive:
            states, golive = [sched[-1]] + live, ts(sched[-1][3])
        first, bare = None, None        # bare: a headline change published with no new note
        shown, taken = {states[0][5]}, set()
        for i in range(1, len(states)):
            a, b = states[i - 1], states[i]
            when = ts(b[1])
            repeat = b[5] in shown      # a note the document had already shown, put back
            shown.add(b[5])
            if not b[5] or b[5] == a[5] or repeat:
                bare = None if b[5] != a[5] else (when if b[4] != a[4] and i not in taken else bare)
                continue
            fixes = sorted(f for f in entry_fix.get(did, ()) if when - dt.timedelta(minutes=20) <= f <= when)
            after = None
            if b[4] == a[4] and not (bare is not None and when - bare <= dt.timedelta(minutes=20)):
                for j in range(i + 1, len(states)):
                    if ts(states[j][1]) - when > dt.timedelta(minutes=20) or states[j][5] != states[j - 1][5]:
                        break
                    if states[j][4] != states[j - 1][4]:
                        after = j
                        break
            # a note on an unchanged headline records the bare headline change just before it, or just after it, or
            # an entry's fix
            if b[4] != a[4]:
                at = when
            elif bare is not None and when - bare <= dt.timedelta(minutes=20):
                at = bare
            elif after is not None:
                at = ts(states[after][1])
                taken.add(after)
            else:
                at = fixes[0] if fixes else None
            bare = None
            if at is not None:
                count[desk] += 1
                march[0] += m_lo <= at < m_hi
                first = at if first is None or at < first else first
            else:
                march[1] += m_lo <= when < m_hi
        if first is not None:
            mins[desk].append(int((first - golive).total_seconds() // 60))
    for d in SHORTLIST:
        check("articles %s equal the spine's base-year articles" % d, arts[d] == C[d]["n_win"], "%d %d" % (arts[d], C[d]["n_win"]))
        check("headline corrections %s" % d, count[d] == CLAIMS["corrections"][d], str(count[d]))
        check("rate per 1,000 articles %s" % d, one_dp(1000 * count[d] / arts[d]) == CLAIMS["rate"][d],
              "%.4f" % (1000 * count[d] / arts[d]))
        check("median minutes to first headline correction %s" % d,
              len(mins[d]) % 2 == 1 and statistics.median(mins[d]) == CLAIMS["median_minutes"][d],
              "%s over %d articles" % (statistics.median(mins[d]), len(mins[d])))
    bull = find("standards_bulletin")
    with zipfile.ZipFile(bull) as z:
        text = re.sub(r"<[^>]+>", " ", z.read("word/document.xml").decode("utf-8"))
    mm = re.search(r"logged (\d+) headline corrections and (\d+) corrections to article text", re.sub(r"\s+", " ", text))
    check("the bulletin's March figures reproduce under the standards rule",
          mm is not None and (int(mm.group(1)), int(mm.group(2))) == tuple(march) == CLAIMS["bulletin_march"],
          "%s %s" % (mm.groups() if mm else None, march))

    n_ok = sum(1 for r in results if r[1])
    for name, ok, det in results:
        print("%s  %s%s" % ("PASS" if ok else "FAIL", name, ("  [%s]" % det) if det else ""))
    print("\nverifier: %d checks, %d passed, %d failed" % (len(results), n_ok, len(results) - n_ok))
    return 0 if n_ok == len(results) else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else str(Path(__file__).resolve().parent.parent / "target")))
