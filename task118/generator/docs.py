"""In-fiction documents for the task118 pack: texts and their PDF, DOCX, EML, MD and TXT writers."""
import datetime as dt
import html
import io
import math
import re
import zipfile

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P

MID = "·"


# ---------------------------------------------------------------------------------- writers

def _styles():
    ss = getSampleStyleSheet()
    base = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=9.6, leading=13.2,
                          spaceAfter=5, alignment=TA_LEFT)
    return {
        "title": ParagraphStyle("t", parent=base, fontName="Helvetica-Bold", fontSize=15, leading=19, spaceAfter=4),
        "meta": ParagraphStyle("m", parent=base, fontSize=8.4, leading=11, textColor=colors.HexColor("#444444"),
                               spaceAfter=2),
        "h": ParagraphStyle("h", parent=base, fontName="Helvetica-Bold", fontSize=10.6, leading=14, spaceBefore=8,
                            spaceAfter=3),
        "p": base,
        "small": ParagraphStyle("s", parent=base, fontSize=8, leading=10.5, textColor=colors.HexColor("#555555")),
        "cell": ParagraphStyle("c", parent=base, fontSize=8.6, leading=11, spaceAfter=0),
    }


# The application each in-fiction PDF was written with: (Creator, Producer, PDF version it writes).
PDF_APPS = {"word2010": ("Microsoft\xae Word 2010", "Microsoft\xae Word 2010", b"1.5"),
            "pdfmaker": ("Acrobat PDFMaker 15 for Word", "Adobe PDF Library 15.0", b"1.6")}


def pdf_literal(s):
    """A PDF literal string in PDFDocEncoding (Latin-1 for every character used here), escaped."""
    out = b""
    for ch in s.encode("latin-1"):
        c = bytes([ch])
        if c in b"()\\":
            out += b"\\" + c
        elif ch > 126 or ch < 32:
            out += b"\\%03o" % ch
        else:
            out += c
    return b"(" + out + b")"


def finish_pdf(path, when, author, title, app):
    """Give the file the information dictionary and header its application writes, drop the writing library's
    comments, and rebuild the cross-reference offsets. Asserts the result parses strictly with unchanged text."""
    import logging
    from pypdf import PdfReader
    creator, producer, version = PDF_APPS[app]
    before = [p.extract_text() for p in PdfReader(str(path)).pages]
    b = open(path, "rb").read()
    head = re.match(rb"%PDF-1\.\d\n%[^\n]*\n", b)
    b = b"%PDF-" + version + b"\n%\xe2\xe3\xcf\xd3\n" + b[head.end():]
    info_no = int(re.search(rb"/Info (\d+) 0 R", b).group(1))
    stamp = when.strftime("D:%Y%m%d%H%M%S").encode() + b"+10'00'"
    info = (b"<<\n/Author " + pdf_literal(author) + b" /CreationDate (" + stamp + b") /Creator " + pdf_literal(creator)
            + b" /ModDate (" + stamp + b") /Producer " + pdf_literal(producer) + b" /Title " + pdf_literal(title)
            + b"\n>>")
    m = re.search(rb"\n%d 0 obj\n(<<.*?\n>>)\nendobj" % info_no, b, re.S)
    b = b[:m.start(1)] + info + b[m.end(1):]
    b = re.sub(rb"\n%[^\n]*ReportLab[^\n]*\n", b"\n", b)
    xref_at = b.rfind(b"\nxref\n") + 1
    hdr = re.match(rb"xref\n0 (\d+)\n", b[xref_at:])
    n = int(hdr.group(1))
    offs, pos = [], 0
    for k in range(1, n):
        mo = re.compile(rb"\n%d 0 obj\n" % k).search(b, pos)
        offs.append(mo.start() + 1)
        pos = mo.end()
    table = hdr.group(0) + b"0000000000 65535 f \n" + b"".join(b"%010d 00000 n \n" % o for o in offs)
    b = b[:xref_at] + table + b[xref_at + len(hdr.group(0)) + 20 * n:]
    b = re.sub(rb"startxref\n\d+\n", b"startxref\n%d\n" % xref_at, b)
    assert b"reportlab" not in b.lower() and b"ReportLab" not in b
    open(path, "wb").write(b)
    cap = io.StringIO()
    h = logging.StreamHandler(cap)
    logging.getLogger("pypdf").addHandler(h)
    try:
        r = PdfReader(str(path), strict=True)
        after = [p.extract_text() for p in r.pages]
        meta = r.metadata
    finally:
        logging.getLogger("pypdf").removeHandler(h)
    assert not cap.getvalue().strip(), "%s: %s" % (path, cap.getvalue())
    assert after == before, path
    assert meta.producer == producer and meta.creator == creator and meta.author == author, meta


def write_pdf(path, title, blocks, footer, when, author="Bightline News", app="word2010"):
    """blocks: list of (kind, payload): title/meta/h/p/small/bullets/table/space."""
    st = _styles()
    flow = []
    for kind, payload in blocks:
        if kind in ("title", "meta", "h", "p", "small"):
            flow.append(Paragraph(payload, st[kind]))
        elif kind == "bullets":
            for b in payload:
                flow.append(Paragraph(b, st["p"], bulletText="•"))
        elif kind == "space":
            flow.append(Spacer(1, payload * mm))
        elif kind == "table":
            rows, widths = payload
            data = [[Paragraph(str(c), st["cell"]) for c in r] for r in rows]
            t = Table(data, colWidths=[w * mm for w in widths], hAlign="LEFT")
            t.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#333333")),
                ("LINEBELOW", (0, -1), (-1, -1), 0.4, colors.HexColor("#999999")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
            ]))
            flow.append(t)
            flow.append(Spacer(1, 3 * mm))

    def on_page(canv, doc):
        canv.saveState()
        canv.setFont("Helvetica", 7.5)
        canv.setFillColor(colors.HexColor("#555555"))
        canv.drawString(18 * mm, 10 * mm, footer)
        canv.drawRightString(A4[0] - 18 * mm, 10 * mm, "Page %d" % doc.page)
        canv.restoreState()

    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title=title, author=author, subject=title,
                            creator=author, producer=author, keywords="", invariant=1)
    doc.build(flow, onFirstPage=on_page, onLaterPages=on_page)
    # invariant mode stamps 2000-01-01 and the library's own name; give the file its application's dictionary
    finish_pdf(path, when, author, title, app)


def write_docx(path, paras, when, author="Bightline News", title=""):
    import docx
    from docx.shared import Pt
    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for kind, text in paras:
        if kind == "title":
            d.add_heading(text, level=1)
        elif kind == "h":
            d.add_heading(text, level=2)
        elif kind == "p":
            d.add_paragraph(text)
        elif kind == "small":
            p = d.add_paragraph()
            r = p.add_run(text)
            r.font.size = Pt(8.5)
    cp = d.core_properties
    cp.author = author
    cp.last_modified_by = author
    cp.title = title
    cp.comments = ""
    cp.subject = ""
    cp.keywords = ""
    cp.category = ""
    cp.revision = 3
    cp.created = when - P.AEST                      # Word records UTC
    cp.modified = when + dt.timedelta(hours=2) - P.AEST
    d.save(str(path))
    repack_ooxml(path, when)


def word_container(data, pages):
    """Word's own statistics in docProps/app.xml, counted from the body text, and the template's thumbnail dropped."""
    body = data["word/document.xml"].decode("utf-8")
    paras = []
    for p in re.findall(r"<w:p[ >].*?</w:p>", body, flags=re.S):
        txt = html.unescape("".join(re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", p)))
        if txt.strip():
            paras.append(txt)
    stats = (("Pages", pages), ("Words", sum(len(p.split()) for p in paras)),
             ("Characters", sum(len(re.sub(r"\s", "", p)) for p in paras)),
             ("Lines", sum(max(1, math.ceil(len(p) / 92)) for p in paras)), ("Paragraphs", len(paras)),
             ("CharactersWithSpaces", sum(len(p) for p in paras)))
    app = data["docProps/app.xml"].decode("utf-8")
    for tag, val in stats:
        app = re.sub(r"<%s>\d+</%s>" % (tag, tag), "<%s>%d</%s>" % (tag, val, tag), app)
    data["docProps/app.xml"] = app.encode("utf-8")
    data.pop("docProps/thumbnail.jpeg", None)
    data["_rels/.rels"] = re.sub(rb'<Relationship [^>]*Target="docProps/thumbnail.jpeg"/>', b"", data["_rels/.rels"])
    if not any(n.endswith((".jpeg", ".jpg")) for n in data):
        data["[Content_Types].xml"] = data["[Content_Types].xml"].replace(
            b'<Default Extension="jpeg" ContentType="image/jpeg"/>', b"")


def repack_ooxml(path, when, pages=1):
    """Rewrite an OOXML zip with fixed entry times, the in-fiction application name and Word's own statistics."""
    zin = zipfile.ZipFile(path)
    infos = zin.infolist()
    data = {i.filename: zin.read(i.filename) for i in infos}
    zin.close()
    if "docProps/app.xml" in data:
        data["docProps/app.xml"] = re.sub(rb"<Application>[^<]*</Application>", b"<Application>Microsoft Office Word</Application>",
                                          data["docProps/app.xml"])
    if "docProps/core.xml" in data:
        data["docProps/core.xml"] = re.sub(rb"<dc:description>[^<]*</dc:description>", b"<dc:description></dc:description>",
                                           data["docProps/core.xml"])
    if "word/document.xml" in data:
        word_container(data, pages)
    infos = [i for i in infos if i.filename in data]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for i in infos:
            zi = zipfile.ZipInfo(i.filename, date_time=(1980, 1, 1, 0, 0, 0))     # the entry time Word writes
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            zi.create_system = 0
            zout.writestr(zi, data[i.filename])


# ---------------------------------------------------------------------------------- the charter (governing)

def charter_blocks():
    rows = [["Desk", "Desk code", "Vertical", "Edition"]]
    for code in P.SHORTLIST:
        d = P.DESK[code]
        rows.append([d[1], code, d[2], "Brisbane" if d[3] == "Brisbane" else "National"])
    return [
        ("title", "Headline squad: 2027 placement brief"),
        ("meta", "Bightline News | Newsroom planning | For the planning meeting on Wednesday 4 November 2026"),
        ("meta", "Owner: Corey Cox, Managing editor | Version 2, 12 October 2026 | Internal"),
        ("space", 3),
        ("h", "What the squad is for"),
        ("p", "The headline squad is six people: a squad lead, three headline editors, a producer and a data analyst. "
              "Each January it moves to one desk for the calendar year, tests that desk's headlines and ships each "
              "winning headline. It takes the headlines the desk has not put under test itself; every headline a desk "
              "or the squad puts under test is recorded in the test archive when the test concludes, whether or not a "
              "variant ships."),
        ("h", "How the squad is judged"),
        ("p", "The squad is judged the way Experimentation closes each embedding: on the article clicks Bightline's "
              "readers make in its year that they would not have made without it, measured against a reader holdback "
              "that keeps every original headline. Clicks count from each article's publication, the hours its "
              "headline is under test included. Clicks in those hours count at the winning headline's lift. "
              "Experimentation scores each test by its winning lift: the shipped headline's click-through over the "
              "control's, minus one, and nil where the control is kept."),
        ("p", "A shipped headline replaces the old one wherever Bightline sets the headline: the home pages, the section "
              "fronts, the app home feed and tabs, and the related-story links on article pages. Search results, "
              "Discover-style cards, partner apps, social link cards, newsletters and alerts keep the headline stored at "
              "publication."),
        ("p", "The 2027 audience plan holds every desk at its clicks in the twelve months to September 2026. New "
              "initiatives for 2027, the squad's year included, are planned on top of those figures."),
        ("h", "Where it has been"),
        ("p", "Since 2019 the squad has worked with the app desks: Games, Wellness, Puzzles and Recipes. Its 2026 "
              "year with Puzzles finishes on Friday 18 December. Experimentation keeps the change log of closed "
              "embeddings."),
        ("h", "2027 shortlist"),
        ("table", (rows, [52, 26, 30, 30])),
        ("p", "The shortlist was agreed at the September editors' meeting. Politics" + MID + "metro and the app desks "
              "are not being considered for 2027."),
        ("h", "How the decision is made"),
        ("p", "The planning meeting takes one recommendation from the managing editor: the desk, and the clicks the "
              "squad is expected to add there in 2027."),
        ("p", "Papers: a short paper and its workbook to Corey Cox by Friday 30 October, for circulation with the "
              "agenda on Monday 2 November."),
        ("h", "Moving in January"),
        ("bullets", [
            "The squad moves on Monday 11 January 2027 and sits with the desk for the year.",
            "The desk editor names a liaison editor by Tuesday 1 December.",
            "Weekly stand-up with the desk editor; monthly report to the managing editor.",
            "Desk access and hardware requests go to the newsroom systems team by Friday 4 December.",
        ]),
        ("space", 2),
        ("small", "Distribution: Lisa Jennings, Kayla Torres, Nina Franklin, Willie Smith; editors of the six "
                  "shortlisted desks."),
    ]


def policy_blocks():
    return [
        ("title", "Corrections"),
        ("meta", "Bightline News editorial standards, section 7 | Version 3.3, effective 1 October 2025, amended 27 November 2025 | "
                 "Owner: Standards editor"),
        ("space", 3),
        ("h", "7.1 Principle"),
        ("p", "We correct errors of fact promptly, openly and where the error appeared. A correction says what was "
              "wrong in plain words. We do not quietly change a published fact."),
        ("h", "7.2 What needs a correction"),
        ("p", "Any error of fact in a headline, standfirst, body text, caption, chart or byline, once it is "
              "confirmed by the desk editor or the duty editor. Changes of style, spelling of common words or "
              "updates to a developing story are not corrections."),
        ("h", "7.3 How a correction is made"),
        ("p", "The editor making the correction fixes the error and adds a correction note in the CMS, on the same "
              "save or on a save no more than ten minutes before or after it. An autosave of an editor's working copy "
              "publishes nothing: readers see the fix and the note once the editor saves them. The note appears at the foot of the "
              "article and stays on every later version until an editor clears it; where a save drops a note and a "
              "later save puts the same note back, that is the original correction, not a new one. Where a second "
              "error is found later, a further note is added above the first."),
        ("h", "7.4 Headlines"),
        ("p", "A headline correction is logged on the version that publishes the corrected headline. Time to "
              "correction runs from when the article went live: when it was first published, by a live save or, for "
              "a scheduled article, when the CMS published it at its scheduled time."),
        ("h", "7.5 Live blogs"),
        ("p", "An error in a live blog entry is fixed in the entry and the correction note is added to the live "
              "blog. The correction is the live blog's: an entry is part of its live blog, not an article of its own."),
        ("h", "7.6 Speed"),
        ("p", "Headline errors are fixed as soon as they are confirmed. Duty editors aim to confirm a reported "
              "headline error within the hour."),
        ("h", "7.7 Removing a note"),
        ("p", "Correction notes stay with an article. The standards editor can agree to remove a note, for example "
              "after a legal settlement or where the note itself was wrong."),
        ("h", "7.8 Reporting"),
        ("p", "The standards desk reports corrections to the editor-in-chief each month in the standards bulletin, "
              "and to the board each quarter. Articles moved across from the previous CMS at the 1 October 2025 "
              "migration went live in that system, so their correction times are not reported from the new CMS. "
              "Where the Brisbane instance restore of 14 November 2025 recreated a document, the recreated copy and "
              "the document it was recreated from are one article."),
    ]


def bulletin_paras(h, b):
    return [
        ("title", "Standards bulletin"),
        ("small", "April 2026 | Issue 31 | From the standards desk to the editor-in-chief and desk editors"),
        ("h", "March in numbers"),
        ("p", "The web desks logged %d headline corrections and %d corrections to article text in March. Both "
              "figures cover the seven web desks, national and Brisbane." % (h, b)),
        ("h", "Titles and names"),
        ("p", "Several of March's corrections were to a person's title or role. The house style list of current "
              "titles was updated on 2 April and is linked from the CMS help page."),
        ("h", "Readers' editor"),
        ("p", "The quarterly readers' editor report is due on 20 April. Desk editors will get their section a week "
              "before."),
        ("small", "Internal. Standards desk."),
    ]


# ---------------------------------------------------------------------------------- the syndication agreement

def agreement_blocks():
    return [
        ("title", "Content syndication agreement"),
        ("meta", "Between Bightline News Pty Ltd (Bightline) and Newsfold Pty Ltd (Newsfold)"),
        ("meta", "Dated 1 July 2024 | Variation 1, 1 July 2025"),
        ("space", 3),
        ("h", "1. Definitions"),
        ("p", "<b>Articles</b> means the articles Bightline publishes on its national edition, and Brisbane edition "
              "sport and politics articles from 1 July 2025 (Variation 1). <b>Feed</b> means the structured feed of "
              "Articles described in Schedule 1. <b>Partner App</b> means the Newsfold news app on iOS, Android and "
              "the web."),
        ("h", "2. Licence"),
        ("p", "Bightline grants Newsfold a non-exclusive licence to display Articles from the Feed in the Partner "
              "App in Australia and New Zealand for the Term. Newsfold may not edit an Article, except to fit its "
              "layout, and may not display an Article after Bightline withdraws it from the Feed."),
        ("h", "3. Exclusions"),
        ("p", "Brisbane edition local news, app-only products (including Games, Puzzles, Recipes and Wellness) and "
              "any Article marked for subscribers only are excluded from the Feed."),
        ("h", "4. Attribution"),
        ("p", "Each Article shows the Bightline masthead and byline and links to the Article on Bightline's site."),
        ("h", "5. Commercial terms"),
        ("p", "Newsfold pays Bightline 50 per cent of net advertising revenue earned on Article pages in the Partner "
              "App, with a minimum guarantee of A$185,000 a year, paid quarterly in arrears."),
        ("h", "6. Reporting"),
        ("p", "Newsfold provides a monthly report of Article views and outbound clicks to Bightline's site, by "
              "Article, within ten business days of month end."),
        ("h", "7. Term and termination"),
        ("p", "The Term is three years from 1 July 2024. Either party may end the agreement on 90 days' written "
              "notice after 30 June 2026."),
        ("h", "8. General"),
        ("p", "This agreement is governed by the law of New South Wales. Each party keeps the terms confidential."),
        ("h", "Schedule 1: the Feed"),
        ("p", "Bightline makes the Feed available at an address it notifies to Newsfold. Newsfold polls the Feed at "
              "least every five minutes and removes withdrawn Articles within one hour."),
        ("space", 4),
        ("small", "Signed for Bightline News Pty Ltd by Lisa Jennings, Editor-in-chief, under delegation. "
                  "Signed for Newsfold Pty Ltd by its director."),
    ]


# ---------------------------------------------------------------------------------- the planning thread

def thread_eml():
    hdr = """From: Corey Cox <corey.cox@bightline.com.au>
To: Kayla Torres <kayla.torres@bightline.com.au>, Nina Franklin <nina.franklin@bightline.com.au>,
 Willie Smith <willie.smith@bightline.com.au>, Natalie Benjamin <natalie.benjamin@bightline.com.au>
Cc: Lisa Jennings <lisa.jennings@bightline.com.au>
Subject: Re: Squad placement 2027
Date: Fri, 16 Oct 2026 08:38:10 +1000
Message-ID: <CAB7q2kx4Lr8T@mail.bightline.com.au>
In-Reply-To: <CAB7q2kx3Vn1P@mail.bightline.com.au>
MIME-Version: 1.0
Content-Type: text/plain; charset="utf-8"
Content-Transfer-Encoding: 8bit

"""
    body = """Thanks all. Natalie, that folder is what I'll work from. I'll take one recommendation to the 4 November
meeting with a number against it.

Corey Cox
Managing editor

On Fri, 16 Oct 2026 at 08:05, Natalie Benjamin <natalie.benjamin@bightline.com.au> wrote:
> Exports are in the 2027 planning folder: pageviews by source and age band for October to September,
> the page analytics monthly export for the same months, the full headline test archive (app engine
> since 2019, web since the migration), the staff list as at Monday, the 2027 plan, the squad change
> log, the panel monthly file with its reference workbook, and the CMS revision export for the web desks
> that Standards asked for last month. The pageview extract is article pageviews only, so home pages
> and section fronts are not in it.
>
> Natalie
>
> On Wed, 14 Oct 2026 at 11:20, Willie Smith <willie.smith@bightline.com.au> wrote:
>> The squad can move on 11 January wherever it lands. I'll need the desk's liaison editor named by
>> 1 December so we can book the stand-ups.
>>
>> Willie
>>
>> On Wed, 14 Oct 2026 at 09:02, Nina Franklin <nina.franklin@bightline.com.au> wrote:
>>> The change log is up to date to embedding 7. The Puzzles year goes in once it closes in December.
>>>
>>> Nina
>>>
>>> On Tue, 13 Oct 2026 at 16:48, Kayla Torres <kayla.torres@bightline.com.au> wrote:
>>>> Panel numbers for the year are in the folder too.
>>>>
>>>> Kayla
>>>>
>>>> On Tue, 13 Oct 2026 at 08:31, Corey Cox <corey.cox@bightline.com.au> wrote:
>>>>> We place the squad for 2027 at the 4 November planning meeting. I want a recommendation I can
>>>>> defend: the desk, and what the squad adds there in clicks. Send me whatever you think matters.
>>>>>
>>>>> Corey

--
Bightline News | Internal correspondence. Please do not forward outside the newsroom.
"""
    return hdr + body
