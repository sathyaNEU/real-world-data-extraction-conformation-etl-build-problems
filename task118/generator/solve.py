"""Every owed figure, from the pack alone: the current users, each one's first start (the root of the repeat-template
chain, through Quenford Rise's previous system), the admissions each locality's team would avoid, and the falls
figures from the SUS cohort rebuilt into stays. The wrong routes are kept beside the golden for the checks."""
import collections
import csv
import datetime as dt
import os

from openpyxl import load_workbook

import model as M
import params as P

D = dt.timedelta
EMERGENCY = {"21", "22", "23", "24", "25", "28", "2A", "2B", "2C", "2D"}
TRANSFER_IN = {"2B", "81"}
HOSPITAL_DEST = {"49", "50", "51", "52", "53", "84", "87"}
HOSPITAL_SOURCE = {"49", "50", "51", "52", "53", "87"}


def read(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def day(s):
    return dt.date.fromisoformat(s) if s else None


def ukday(s):
    return dt.datetime.strptime(s, "%d/%m/%Y").date() if s else None


def register(target):
    wb = load_workbook(os.path.join(target, P.FILES["register"]), read_only=True)
    ws = wb["Practices"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    out = [dict(zip(hdr, r)) for r in rows[1:]]
    loc = {r["Practice code"]: r["Locality code"] for r in out}
    for r in out:
        for c in str(r["Previous practice codes"] or "").replace(";", ",").split(","):
            if c.strip():
                loc[c.strip()] = r["Locality code"]
    return out, loc


def current_users(target, loc):
    """aged 65 or over on 1 January 2027 (born 1961 or earlier) with a benzodiazepine or Z-drug dispensed July to
    September 2026, in the locality of their practice"""
    first, inq, yob, where = {}, set(), {}, {}
    for r in read(os.path.join(target, P.FILES["dispensing"])):
        p, d = r["patient_pseudo_id"], day(r["dispensed_date"])
        if p not in first or d < first[p]:
            first[p] = d
        if P.CURRENT_FROM <= d <= P.CURRENT_TO:
            inq.add(p)
            where[p] = loc[r["practice_code"]]
        yob[p] = int(r["year_of_birth"])
    cur = sorted(p for p in inq if yob[p] <= P.AGE_ON.year - 66)
    return cur, {p: where[p] for p in cur}, first


def starts(target):
    """each patient's first start: walk the current system's chain back by replaces_template_id; where the move
    re-created a template, carry on in the previous system's export by SupersededRef; the root's date is the start"""
    tm = read(os.path.join(target, P.FILES["templates"]))
    lg = read(os.path.join(target, P.FILES["legacy"]))
    by_id = {t["template_id"]: t for t in tm}
    leg = {t["TemplateRef"]: t for t in lg}
    latest = {}
    for t in tm:
        p = t["patient_pseudo_id"]
        if p not in latest or t["authorised_from"] > latest[p]["authorised_from"]:
            latest[p] = t
    out, cur_root, latest_from, hop1 = {}, {}, {}, {}
    for p, t in latest.items():
        latest_from[p] = day(t["authorised_from"])
        while t["replaces_template_id"] in by_id:
            t = by_id[t["replaces_template_id"]]
        cur_root[p] = day(t["authorised_from"])
        if t["migration_source_ref"]:
            lt = leg[t["migration_source_ref"]]
            hop1[p] = ukday(lt["StartDate"])
            while lt["SupersededRef"] in leg:
                lt = leg[lt["SupersededRef"]]
            out[p] = ukday(lt["StartDate"])
        else:
            hop1[p] = cur_root[p]
            out[p] = cur_root[p]
    earliest = collections.defaultdict(lambda: dt.date(9999, 1, 1))
    for t in tm:
        if not t["migration_source_ref"]:
            earliest[t["patient_pseudo_id"]] = min(earliest[t["patient_pseudo_id"]], day(t["authorised_from"]))
    for t in lg:
        earliest[t["PatientPseudoID"]] = min(earliest[t["PatientPseudoID"]], ukday(t["StartDate"]))
    return dict(golden=out, current_root=cur_root, current_template=latest_from, one_hop=hop1, earliest=dict(earliest))


def by_locality(users, where, start_of):
    tot = collections.defaultdict(float)
    for p in users:
        tot[where[p]] += M.excess_years(start_of(p))
    return {c: P.STOP * P.BASE / 100.0 * tot.get(c, 0.0) for c, _, _, _ in P.LOCALITIES}


def call(target):
    reg, loc = register(target)
    users, where, first = current_users(target, loc)
    st = starts(target)
    missing = [p for p in users if p not in st["golden"]]
    routes = {
        "golden": by_locality(users, where, lambda p: st["golden"][p]),
        "earliest_both_systems": by_locality(users, where, lambda p: st["earliest"][p]),
        "extract_first_dispensing": by_locality(users, where, lambda p: first[p]),
        "current_template": by_locality(users, where, lambda p: st["current_template"][p]),
        "current_system_root": by_locality(users, where, lambda p: st["current_root"][p]),
        "previous_system_one_hop": by_locality(users, where, lambda p: st["one_hop"][p]),
        "users": {c: float(sum(1 for p in users if where[p] == c)) for c, _, _, _ in P.LOCALITIES},
    }
    for k in (21, 28, 35, 56):
        cut = P.EXTRACT_FROM + D(days=k)
        routes["prevalent_long_term_%d" % k] = by_locality(
            users, where, lambda p, cut=cut: first[p] if first[p] >= cut else dt.date(2020, 1, 1))
    g = routes["golden"]
    order = sorted(g, key=lambda c: -g[c])
    return dict(users=users, where=where, first=first, starts=st, missing=missing, routes=routes, order=order,
                winner=order[0], runner=order[1], avoided=g[order[0]], runner_avoided=g[order[1]],
                n_users={c: sum(1 for p in users if where[p] == c) for c, _, _, _ in P.LOCALITIES})


# ---------------------------------------------------------------------------------------------------- the asks

def spells(target):
    rows = read(os.path.join(target, P.FILES["sus"]))
    sp = {}
    for n, r in enumerate(rows):
        k = (r["provider_code"], r["hospital_spell_no"])
        s = sp.setdefault(k, dict(key=k, pid=r["sus_pseudo_id"], adm=day(r["admission_date"]), row=n,
                                  site=r["site_code"], method=r["admission_method"], source=r["admission_source"],
                                  eps=[]))
        s["eps"].append(r)
    for s in sp.values():
        s["eps"].sort(key=lambda r: int(r["episode_no"]))
        last = s["eps"][-1]
        s["dis"] = day(last["discharge_date"])
        s["dest"] = last["discharge_destination"]
        s["dmethod"] = last["discharge_method"]
        e1 = s["eps"][0]
        s["age"] = int(e1["age_at_start"])
        s["practice"] = e1["gp_practice_code"]
        s["diags"] = [e1["primary_diagnosis"]] + [e1["diag_%02d" % i] for i in range(2, 9)]
    return list(sp.values())


def is_fall(diags):
    return any(c and c[0] == "W" and "00" <= c[1:3] <= "19" for c in diags)


def stays(sp, link="dest", gap=1):
    """join a patient's spells into stays. The golden ("dest"): a spell whose patient was moved to another hospital
    runs on into the patient's spell at another site beginning the day they left or the next, however the receiving
    hospital coded it. The other modes are the routes a careful reader takes instead (checks and pair.py)."""
    by_p = collections.defaultdict(list)
    for s in sp:
        by_p[s["pid"]].append(s)
    out = []
    for pid, ss in by_p.items():
        if link == "shift":
            ss.sort(key=lambda s: (s["adm"], s["row"]))
        else:
            ss.sort(key=lambda s: (s["adm"], s["dis"], s["key"]))
        succ, taken = {}, set()
        for i, a in enumerate(ss):
            if link == "none":
                break
            if link == "shift":
                b = ss[i + 1] if i + 1 < len(ss) else None
                if b is not None and a["dest"] in HOSPITAL_DEST and 0 <= (b["adm"] - a["dis"]).days <= gap \
                        and id(b) not in taken:
                    succ[id(a)] = b
                    taken.add(id(b))
                continue
            for b in ss:
                if b is a or id(b) in taken or b["adm"] < a["dis"] or (b["adm"] == a["adm"] and b["dis"] < a["dis"]):
                    continue
                lag = (b["adm"] - a["dis"]).days
                if lag > gap:
                    continue
                if link == "consecutive":
                    ok = True
                else:
                    ok = a["dest"] in HOSPITAL_DEST and b["site"] != a["site"]
                    if link == "dest_src":
                        ok = ok and b["source"] in HOSPITAL_SOURCE
                    elif link == "dest_method":
                        ok = ok and b["method"] in TRANSFER_IN
                    elif link == "dest_either":
                        ok = ok and (b["source"] in HOSPITAL_SOURCE or b["method"] in TRANSFER_IN)
                if ok:
                    succ[id(a)] = b
                    taken.add(id(b))
                    break
        for s in ss:
            if id(s) in taken:
                continue
            chain = [s]
            while id(chain[-1]) in succ:
                chain.append(succ[id(chain[-1])])
            out.append(dict(pid=pid, spells=chain, start=chain[0]["adm"], end=chain[-1]["dis"],
                            died=chain[-1]["dmethod"] == "4", age=chain[0]["age"], practice=chain[0]["practice"],
                            method=chain[0]["method"], fall=is_fall(chain[0]["diags"])))
    return out


def asks(target, link="dest", gap=1, loc=None, emergency=None, readmit_from="end"):
    if loc is None:
        loc = register(target)[1]
    emergency = emergency or (EMERGENCY - TRANSFER_IN)
    st = stays(spells(target), link, gap)
    by_p = collections.defaultdict(list)
    for s in st:
        by_p[s["pid"]].append(s)
    out = {c: dict(falls_65_79=0, falls_80=0, bed_days=0, readmitted=0) for c, _, _, _ in P.LOCALITIES}
    for s in st:
        if not (s["fall"] and s["method"] in emergency and P.FY_FROM <= s["start"] <= P.FY_TO):
            continue
        c = loc.get(s["practice"])
        if c is None:
            continue
        o = out[c]
        o["falls_65_79" if s["age"] < 80 else "falls_80"] += 1
        o["bed_days"] += (s["end"] - s["start"]).days
        if not s["died"]:
            left = s["end"] if readmit_from == "end" else s["spells"][0]["dis"]
            if any(x is not s and x["method"] in emergency and left <= x["start"] <= left + D(days=30)
                   for x in by_p[s["pid"]]):
                o["readmitted"] += 1
    return out
