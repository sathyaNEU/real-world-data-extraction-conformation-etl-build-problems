"""Independent verifier: reads only the shipped bundle (target.zip unpacked to a scratch folder), on its own code
path (pandas, no generator module), and recomputes every graded figure against figures.json. Reads back every format
the pack ships, with a positive control on each extractor.

    python3 task118/generator/verify.py [path/to/target.zip]
"""
import json
import os
import re
import sys
import tempfile
import zipfile

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
FAILS = []


def check(label, cond, detail=""):
    print("%s  %s%s" % ("PASS" if cond else "FAIL", label, ("  [%s]" % (detail,)) if not cond else ""))
    if not cond:
        FAILS.append(label)


def pdf(path):
    from pypdf import PdfReader
    return " ".join(p.extract_text() or "" for p in PdfReader(path).pages)


def main(zp=os.path.join(TASK, "target.zip")):
    fig = json.load(open(os.path.join(HERE, "figures.json")))
    tmp = tempfile.mkdtemp()
    with zipfile.ZipFile(zp) as z:
        z.extractall(tmp)
    f = lambda pat: os.path.join(tmp, [x for x in os.listdir(tmp) if re.search(pat, x)][0])

    # every format reads back, each through its extractor with a positive control
    plan = pdf(f(r"^ra_plan_2027"))
    check("pdf read-back: the plan's accreditation sentence", "expected load reduction in the system peak hour of the "
          "design year" in re.sub(r"\s+", " ", plan))
    op7 = re.sub(r"\s+", " ", pdf(f(r"^op7_")))
    check("pdf read-back: OP-7's event-day sentence", "weekday whose day-ahead forecast high" in op7 and "110F" in op7)
    hbn = re.sub(r"\s+", " ", pdf(f(r"^hbn_participation")))
    check("pdf read-back: schedule 2's incentive and cap", "USD 15" in hbn and "USD 255" in hbn and "3.0 kW" in hbn)
    import openpyxl
    ws = openpyxl.load_workbook(f(r"^dr_expansion_costs"), data_only=True)["2027 expansion"]
    cost = {r[0]: int(r[3]) for r in ws.iter_rows(values_only=True) if r and r[2] == "Total, all in"}
    budget = [int(r[3]) for r in ws.iter_rows(values_only=True) if r and isinstance(r[0], str)
              and r[0].startswith("2027 expansion line")][0]
    check("xlsx read-back: four programme costs and the budget line", len(cost) == 4 and budget == 3_200_000)
    reg = json.load(open(f(r"^halvern_device_registry")))["devices"]
    check("json read-back: Halvern's registry", len(reg) > 5000 and "install_address" in reg[0])

    # the design peak hour and its run day
    dy = pd.read_csv(f(r"^appendix_b_design"))
    dy["t"] = pd.to_datetime(dy["hour_beginning_mst"])
    peak = dy.loc[dy["forecast_system_load_mw"].idxmax(), "t"]
    daily = dy.groupby(dy["t"].dt.date)["dry_bulb_f"].max()
    k, d = 0, peak.date()
    hol = {pd.Timestamp("2027-07-05").date(), pd.Timestamp("2027-09-06").date()}
    while d in daily.index and daily[d] >= 110 and pd.Timestamp(d).weekday() < 5 and d not in hol:
        k, d = k + 1, d - pd.Timedelta(days=1)
    check("design peak hour", peak.strftime("%Y-%m-%d %H:%M") == fig["design_peak"], peak)
    check("design peak hour's run day", k == fig["design_run_day"] == 4, k)

    # each event date's run day
    log = pd.read_csv(f(r"^dr_event_log"))
    dates = sorted(pd.to_datetime(log["event_date"].unique()))
    s = set(dates)
    run = {}
    for x in dates:
        kk, y = 1, x - pd.Timedelta(days=1)
        while y in s:
            kk, y = kk + 1, y - pd.Timedelta(days=1)
        run[x] = kk
    fourth = sorted(x for x in dates if run[x] == 4)

    # per-event 18:00 reductions per enrolled unit
    cc = pd.read_csv(f(r"^ridgeway_cool_credit"))
    cc["t"] = pd.to_datetime(cc["interval_start_utc"].str.replace("Z", "")) - pd.Timedelta(hours=7)
    cc = cc[cc["t"].dt.hour == 18]
    g = cc.groupby([cc["t"].dt.normalize(), "assigned_group"])["kwh"].mean().unstack()
    red = {"Cool Credit": g["holdout"] - g["treatment"]}
    pt = pd.read_csv(f(r"^ridgeway_peak_time")).query("hour_beginning_mst == '18:00'")
    pt = pt.pivot(index="event_date", columns="group", values="mean_kw")
    red["Peak Time Rewards"] = pd.Series((pt["holdout"] - pt["alerted"]).values, index=pd.to_datetime(pt.index))
    bd = pd.read_csv(f(r"^ridgeway_business")).query("hour_beginning_mst == '18:00'")
    red["Business Demand Partners"] = bd.groupby(pd.to_datetime(bd["event_date"]))["verified_reduction_kw"].mean()
    fl = pd.read_csv(f(r"^brightwell_fleet")).query("hour_beginning_mst == '18:00' and dispatch_setting_kw == 3.0")
    red["Home Battery Network"] = pd.Series((fl["discharge_kw_ac"] / fl["batteries_enrolled"]).values,
                                            index=pd.to_datetime(fl["event_date"]))
    units = {p: budget // cost[p] for p in cost}
    mw = {}
    for p, ser in red.items():
        v = ser[[x for x in fourth if x in ser.index]].mean()
        mw[p] = v * units[p] / 1000.0
        check("%s per-unit fourth-day reduction" % p, abs(v - fig["per_unit_kw_fourth_day"][p]) < 5e-5,
              (v, fig["per_unit_kw_fourth_day"][p]))
        check("%s whole megawatts" % p, int(round(mw[p])) == fig["mw"][p], (mw[p], fig["mw"][p]))
    best = max(mw, key=mw.get)
    second = sorted(mw, key=mw.get)[-2]
    check("the call", best == fig["call"] == "Home Battery Network", best)
    check("the runner-up", second == fig["runner_up"], second)
    avg = {p: red[p].mean() * units[p] / 1000.0 for p in red}
    avg["Home Battery Network"] = None
    check("the average-event route names Cool Credit (the rung the write-up rejects)",
          max((v, p) for p, v in avg.items() if v is not None)[1] == "Cool Credit")

    # ask A: premises online at 16:00 MST on each 2026 event day
    cis = pd.read_csv(f(r"^cool_credit_enrolments"), dtype=str).rename(columns={"thermostat_serial": "tserial"})
    ab = {"W": "WEST", "E": "EAST", "N": "NORTH", "S": "SOUTH", "DR": "DRIVE", "RD": "ROAD", "LN": "LANE",
          "ST": "STREET", "AVE": "AVENUE", "CT": "COURT", "PL": "PLACE", "BLVD": "BOULEVARD", "TRL": "TRAIL"}
    canon = lambda a: " ".join(ab.get(t, t) for t in a.upper().replace("#", " UNIT ").replace(",", " ").split())
    rg = pd.DataFrame(reg)
    rg["key"] = rg["install_address"].map(canon)
    rg["reg"] = pd.to_datetime(rg["registered_at"].str.replace("Z", ""))
    rg["ret"] = pd.to_datetime(rg["retired_at"].fillna("").str.replace("Z", ""), errors="coerce")
    st = pd.read_csv(f(r"^halvern_status_changes"))
    st["t"] = pd.to_datetime(st["changed_at_utc"].str.replace("Z", ""))
    st = st.sort_values(["serial", "t"])
    un = pd.read_csv(f(r"^halvern_unreachable"))
    switch = None
    for day in sorted(un["event_date"].unique()):
        t14 = pd.Timestamp(day) + pd.Timedelta(hours=21)
        last = st[st["t"] <= t14].groupby("serial").tail(1)
        idle = set(last.loc[last["status"] == "IDLE", "serial"])
        if idle and len(idle - set(un.loc[un["event_date"] == day, "serial"])) / len(idle) > 0.5:
            switch = pd.Timestamp(day)
            break
    check("ask A IDLE switch date", switch is not None and switch.strftime("%Y-%m-%d") == "2026-07-07", switch)
    cis["key"] = cis["service_address"].map(canon)
    devs = cis[["premise_id", "tserial"]].rename(columns={"tserial": "serial"})
    by_addr = cis[["premise_id", "key"]].merge(rg[["serial", "key"]], on="key")[["premise_id", "serial"]]
    links = pd.concat([devs, by_addr]).drop_duplicates()
    links = links.merge(rg[["serial", "reg", "ret"]], on="serial")
    for day, want in fig["ask_a"].items():
        dd = pd.Timestamp(day)
        t = dd + pd.Timedelta(hours=23)
        en = cis[(pd.to_datetime(cis["enrolled_on"]) <= dd) & ((cis["unenrolled_on"].isna()) |
                                                                (pd.to_datetime(cis["unenrolled_on"]) > dd))]
        last = st[st["t"] <= t].groupby("serial").tail(1).set_index("serial")
        on_ser = set(last.index[(last["status"] == "ONLINE") | ((last["status"] == "IDLE") & (last["t"] >= switch))])
        ok = links[(links["reg"] <= t) & ((links["ret"].isna()) | (links["ret"] > t)) & links["serial"].isin(on_ser)]
        got = len(set(ok["premise_id"]) & set(en["premise_id"]))
        check("ask A %s" % day, got == want, (got, want))

    # ask B: incentives by event month and premises at the cap
    rc = pd.read_csv(f(r"^brightwell_dispatch_records")).drop_duplicates()
    rc["s"] = pd.to_datetime(rc["dispatch_start_utc"].str.replace("Z", ""))
    rc["h"] = (pd.to_datetime(rc["dispatch_end_utc"].str.replace("Z", "")) - rc["s"]).dt.total_seconds() / 3600
    rc["day"] = (rc["s"] - pd.Timedelta(hours=7)).dt.normalize()
    ir = pd.read_csv(f(r"^battery_interconnection"), dtype=str)
    rc = rc.merge(ir[["battery_serial", "premise_id"]], on="battery_serial", how="left")
    bat = rc.groupby(["premise_id", "battery_serial", "day"])["h"].sum().reset_index()
    q = bat[bat["h"] >= 3.0][["premise_id", "day"]].drop_duplicates().sort_values(["premise_id", "day"])
    q["n"] = q.groupby("premise_id").cumcount() + 1
    earn = int(re.search(r"the premise earns USD (\d+)", hbn).group(1))
    cap = int(re.search(r"at most USD (\d+) in a season", hbn).group(1))
    paid = q[q["n"] <= cap // earn]
    months = paid.groupby(paid["day"].dt.strftime("%Y-%m")).size() * earn
    for m, want in fig["ask_b_months"].items():
        check("ask B %s" % m, int(months.get(m, 0)) == want, (int(months.get(m, 0)), want))
    capped = int((q.groupby("premise_id").size() >= cap // earn).sum())
    check("ask B premises at the cap", capped == fig["ask_b_capped"], (capped, fig["ask_b_capped"]))

    # ask C: nominated kW by month, CBMU meters, the latest row per meter
    nm = pd.read_csv(f(r"^arroyo_nominations"), dtype={"meter_no": str})
    mt = set(pd.read_csv(f(r"^cbmu_commercial_meters"), dtype={"meter_no": str})["meter_no"])
    nm = nm[nm["meter_no"].isin(mt)].sort_values("submitted_at_utc").groupby(["month", "meter_no"]).tail(1)
    tot = nm.groupby("month")["nominated_kw"].sum()
    for m, want in fig["ask_c_months_whole"].items():
        check("ask C %s" % m, int(round(tot[m])) == want, (tot[m], want))
    print("%d failures" % len(FAILS))
    return 1 if FAILS else 0


if __name__ == "__main__":
    sys.exit(main(*sys.argv[1:]))
