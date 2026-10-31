"""Write task118's pack into a folder: every file under target/, in the layout each producing system exports.

    import build; build.build("task118/target")
"""
import csv
import datetime as dt
import io
import json
import os

import openpyxl
from openpyxl.styles import Font

import docs as D
import params as P
import world as W

FILES = {
    "plan": "ra_plan_2027_capacity_accreditation.pdf",
    "design": "appendix_b_design_year_2027.csv",
    "op7": "op7_demand_response_dispatch_rev3.pdf",
    "log": "dr_event_log_2024-2026.csv",
    "weather": "cbx_daily_weather_jun-sep_2024-2026.csv",
    "cc_panel": "ridgeway_cool_credit_panel_2024-2026.csv",
    "ptr": "ridgeway_peak_time_rewards_2024-2026.csv",
    "bdp": "ridgeway_business_demand_partners_2024-2026.csv",
    "fleet": "brightwell_fleet_event_summary_2024-2026.csv",
    "costs": "dr_expansion_costs_2027.xlsx",
    "hbn_terms": "hbn_participation_agreement_schedule2.pdf",
    "cis": "cool_credit_enrolments_cis_2026-10-16.csv",
    "registry": "halvern_device_registry_2026-10-16.json",
    "status": "halvern_status_changes_2026.csv",
    "unreachable": "halvern_unreachable_devices_2026.csv",
    "dispatch": "brightwell_dispatch_records_2026.csv",
    "interconnect": "battery_interconnection_register_2026-10-16.csv",
    "nominations": "arroyo_nominations_2026.csv",
    "meters": "cbmu_commercial_meters_2026-10-16.csv",
    "capacity_prices": "regional_capacity_price_outlook_2025-2027.csv",
    "satisfaction": "dr_customer_survey_2026.xlsx",
}
PACK = list(FILES)
XLSX_STAMP = dt.datetime(2026, 10, 19, 8, 40, 0)


def _iso(t):
    return t.strftime("%Y-%m-%dT%H:%M:%SZ")


def _csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)


def _xlsx(path, sheets, creator):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows, bold in sheets:
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(list(r))
        for i in bold:
            for c in ws[i]:
                c.font = Font(bold=True)
        for col, width in zip("ABCDEFG", (34, 30, 44, 18, 16, 16, 16)):
            ws.column_dimensions[col].width = width
    wb.properties.creator = creator
    wb.properties.lastModifiedBy = creator
    wb.properties.created = XLSX_STAMP
    wb.properties.modified = XLSX_STAMP
    wb.properties.title = None
    buf = io.BytesIO()
    wb.save(buf)
    # openpyxl writes the zip members with the current time; rewrite them at a fixed stamp for identical rebuilds
    import zipfile
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            data = src.read(info.filename)
            zi = zipfile.ZipInfo(info.filename, date_time=XLSX_STAMP.timetuple()[:6])
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, data)


def build(folder):
    os.makedirs(folder, exist_ok=True)
    f = lambda k: os.path.join(folder, FILES[k])
    wx = W.weather()
    evs = [w for w in wx if w["event"]]

    D.plan(f("plan"))
    D.op7(f("op7"))
    D.hbn(f("hbn_terms"))
    for k, when in (("plan", b"20261012091500"), ("op7", b"20240424143000"), ("hbn_terms", b"20250314110500")):
        raw = open(f(k), "rb").read()
        assert raw.count(b"D:20000101000000") == 2
        open(f(k), "wb").write(raw.replace(b"D:20000101000000", b"D:" + when))

    dy = W.design_year()
    _csv(f("design"), ["hour_beginning_mst", "dry_bulb_f", "forecast_system_load_mw"],
         [(r["ts"].strftime("%Y-%m-%d %H:%M"), "%.1f" % r["temp"], r["load"]) for r in dy])

    _csv(f("weather"), ["STATION", "NAME", "DATE", "TMAX", "TMIN"],
         [("USW00093118", "CHOLLA BASIN INTERNATIONAL AIRPORT, AZ US", w["date"].isoformat(), "%.0f" % w["high"],
           "%.0f" % w["low"]) for w in wx])

    log = W.event_log(evs)
    _csv(f("log"), ["ticket", "programme", "event_date", "window_mst", "forecast_high_f", "notice_issued_mst"],
         [(x["ticket"], x["programme"], x["date"].isoformat(), "16:00-20:00", x["forecast"],
           x["notice"].strftime("%Y-%m-%d %H:%M")) for x in log])

    _csv(f("cc_panel"), ["season", "premise_key", "assigned_group", "interval_start_utc", "kwh"],
         [(s, k, g, _iso(t), "%.3f" % v) for (s, k, g, t, v) in W.cc_panel(evs)])

    _csv(f("ptr"), ["season", "event_date", "hour_beginning_mst", "group", "customers", "mean_kw"],
         [(s, d.isoformat(), "%02d:00" % h, g, n, "%.4f" % m) for (s, d, h, g, n, m) in W.ptr_means(evs)])

    _csv(f("bdp"), ["season", "event_date", "site_key", "hour_beginning_mst", "verified_reduction_kw"],
         [(s, d.isoformat(), k, "%02d:00" % h, "%.1f" % v) for (s, d, k, h, v) in W.bdp_verified(evs)])

    _csv(f("fleet"), ["event_date", "hour_beginning_mst", "batteries_enrolled", "batteries_dispatching",
                      "dispatch_setting_kw", "discharge_kw_ac"],
         [(d.isoformat(), "%02d:00" % h, e, n, "%.1f" % s, "%.1f" % kw) for (d, h, e, n, s, kw) in W.hbn_fleet(evs)])

    cost_rows = [("2027 demand response expansion: cost per added enrolled unit", "", "", ""),
                 ("Budget and Rates, for the 2027 budget round. Updated 2 October 2026.", "", "", ""),
                 ("", "", "", ""),
                 ("Programme", "Unit added", "Cost component", "USD per added unit")]
    for p in P.PROGRAMS:
        for comp, usd in P.COST_PARTS[p]:
            cost_rows.append((P.NAME[p], P.UNIT[p], comp, usd))
        cost_rows.append((P.NAME[p], P.UNIT[p], "Total, all in", P.COST[p]))
    cost_rows += [("", "", "", ""), ("2027 expansion line in the draft budget", "", "", P.BUDGET_USD)]
    _xlsx(f("costs"), [("2027 expansion", cost_rows, [4])], "CBMU Budget and Rates")

    th = W.thermostats()
    _csv(f("cis"), ["premise_id", "service_address", "account_no", "thermostat_serial", "enrolled_on",
                    "unenrolled_on"],
         [(c["premise_id"], c["service_address"], c["account_no"], c["enrolled_serial"], c["enrolled_on"].isoformat(),
           c["unenrolled_on"].isoformat() if c["unenrolled_on"] else "") for c in th["cis"]])
    reg = []
    for d in sorted(th["devices"], key=lambda d: d["registered"]):
        fw = "6.0.%d" % (3 if d["retired"] is None else 2) if (d["retired"] is None or d["retired"].date() >= P.HALVERN_IDLE_SWITCH) \
            else ("5.8.1" if d["model"] == "HC-100" else "5.9.4")
        reg.append({"serial": d["serial"], "model": d["model"], "install_address": d["addr"],
                    "registered_at": _iso(d["registered"]), "retired_at": _iso(d["retired"]) if d["retired"] else None,
                    "firmware": fw})
    with open(f("registry"), "w", encoding="utf-8") as fh:
        json.dump({"account": "Cholla Basin Municipal Utility (Cool Credit)", "exported_at": "2026-10-16T14:02:11Z",
                   "devices": reg}, fh, indent=1)
        fh.write("\n")
    _csv(f("status"), ["serial", "status", "changed_at_utc"],
         [(s, st, _iso(t)) for (s, st, t) in sorted(th["status"], key=lambda x: (x[2], x[0]))])
    un, _hold = W.cc_acks(th, evs)
    _csv(f("unreachable"), ["event_date", "serial", "command", "logged_at_utc"],
         [(d.isoformat(), s, "PRECOOL", _iso(t)) for (d, s, t) in sorted(un, key=lambda x: (x[2], x[1]))])

    ireg, recs = W.batteries(evs)
    _csv(f("interconnect"), ["premise_id", "meter_no", "battery_serial", "inverter_make", "usable_kwh",
                             "hbn_enrolment_no", "energised_on"],
         [(b["premise"], b["meter"], b["serial"], b["make"], "%.1f" % b["usable_kwh"], b["enrolment"],
           b["energised"].isoformat()) for b in sorted(ireg, key=lambda b: (b["energised"], b["serial"]))])
    rows = [("BW%d" % x["rid"], x["enrolment"], x["serial"], _iso(x["start"]), _iso(x["end"]), "%.2f" % x["kwh"],
             x["settled"].isoformat()) for x in recs]
    resent = [r for r, x in zip(rows, recs) if x["settled"] == dt.date(2026, 8, 6)]
    _csv(f("dispatch"), ["record_id", "enrolment_no", "battery_serial", "dispatch_start_utc", "dispatch_end_utc",
                         "delivered_kwh", "settled_on"], rows + resent)

    meters, noms = W.nominations()
    _csv(f("meters"), ["meter_no", "premise_id", "rate_class", "service_address"],
         [(m["meter"], m["premise"], m["rate"], m["addr"]) for m in sorted(meters, key=lambda m: m["premise"])])
    _csv(f("nominations"), ["month", "arroyo_site_ref", "site_name", "service_address", "meter_no", "nominated_kw",
                            "submitted_at_utc"],
         [("2026-%02d" % x["month"], x["ref"], x["name"], x["addr"], x["meter"], "%.1f" % x["kw"], _iso(x["sub"]))
          for x in sorted(noms, key=lambda x: (x["month"], x["sub"]))])

    _csv(f("capacity_prices"), ["year", "month", "bilateral_on_peak_usd_per_kw_month", "bilateral_off_peak_usd_per_kw_month"],
         W.capacity_prices())
    sv = [("2026 demand response customer survey", "", "", ""), ("CBMU Customer Programs, fielded 12 to 30 September 2026", "", "", ""),
          ("", "", "", ""), ("Programme", "Question", "Respondents", "Mean score")] + W.satisfaction()
    _xlsx(f("satisfaction"), [("Results", sv, [4])], "CBMU Customer Programs")
    scrub(folder)
    return evs


def scrub(folder):
    """House fix H1: no writing library's signature or build-clock date in any shipped container."""
    import subprocess
    import sys
    here = os.path.dirname(os.path.abspath(__file__))
    script = os.path.join(here, "..", "..", ".claude", "skills", "reduce-house-fixes", "scripts",
                          "scrub_producer_metadata.py")
    r = subprocess.run([sys.executable, script, folder, "--apply", "--producer", "CBMU", "--stamp",
                        P.EXPORT_DATE.isoformat()], capture_output=True, text=True)
    assert r.returncode == 0, r.stdout + r.stderr
