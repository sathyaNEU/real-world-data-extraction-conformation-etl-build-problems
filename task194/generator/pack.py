"""Write the bundle the network's clinical information unit would hand over."""
import csv
import datetime as dt
import os

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import sim
import world as W

FILES = {
    "extract": "episode_outcomes_extract.csv",
    "detail": "hac_detail_extract.csv",
    "series": "unit_indicator_series.csv",
    "summary": "activity_summary_monthly.csv",
    "plan": "activity_plan_2027q4.csv",
    "units": "unit_register.csv",
    "sites": "hospital_sites.csv",
    "bands": "drg_risk_bands.csv",
    "strata": "hac_stratum_rates.csv",
    "catlist": "hac_category_list.csv",
    "placements": "improvement_placements.csv",
    "dashboard": "quality_dashboard_2026q2.xlsx",
    "dictionary": "data_dictionary.csv",
    "manifest": "export_manifest.csv",
    "experience": "patient_experience_unit_quarterly.csv",
    "nursing": "nursing_hours_by_unit_monthly.csv",
    "hygiene": "hand_hygiene_audit_2026.csv",
}


def _csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        for r in rows:
            w.writerow(r)
    return path


def facility(unit_code, day):
    """The hospital site a unit reported under on a day: three cost centres were realigned on 1 July 2026."""
    if unit_code in W.REALIGNED:
        was, now = W.REALIGNED[unit_code]
        return now if day >= W.REALIGNMENT_DATE else was
    for u in W.UNITS:
        if u[0] == unit_code:
            return u[2]
    raise KeyError(unit_code)


def patient_uid(row, i):
    r = sim.rng("patient", row["unit"], row["sep_date"].isoformat(), i // 3)
    return "MRN%07d" % (1000000 + int(r.integers(0, 850000)))


def write_extract(target, rows, replays):
    """One row per episode record the warehouse holds: every separation from 1 July 2024 to 31 December 2026.

    Episode identifiers run in one network-wide sequence by separation month, so the records re-sent by the
    sub-acute feed take identifiers from the month they belong to, beside the episodes loaded first.
    """
    for r in rows:
        r.setdefault("replay_of", None)
    allrows = list(rows) + list(replays)
    allrows.sort(key=lambda r: (r["sep_date"], r["unit"], r["replay_of"] is not None))
    seq = {}
    out = []
    for i, r in enumerate(allrows):
        key = (r["sep_date"].year, r["sep_date"].month)
        seq[key] = seq.get(key, 0) + 1
        eid = "MR%02d%02d%05d" % (r["sep_date"].year % 100, r["sep_date"].month, seq[key])
        codes = r.get("hac_codes") or []
        out.append([eid, r["unit"], r["pid"], r["adm_date"].isoformat(), r["sep_date"].isoformat(),
                    r["care"], r["los"], r["drg"], r["age"],
                    "" if r["comorbidity"] is None else r["comorbidity"],
                    1 if r["hac"] else 0, codes[0] if codes else ""])
        r["episode_id"] = eid
    _csv(os.path.join(target, FILES["extract"]),
         ["episode_id", "unit_code", "patient_uid", "admission_date", "separation_date", "care_type",
          "los_days", "drg_code", "age_band", "comorbidity_score", "hac_recorded", "hac_category_code"], out)
    return allrows


def write_detail(target, rows, detail):
    out = []
    for d in detail:
        r = rows[d["row"]]
        out.append([r["episode_id"], d["unit"], r["sep_date"].isoformat(), d["seq"], d["code"],
                    d["onset"].isoformat()])
    out.sort(key=lambda x: (x[2], x[0], x[3]))
    return _csv(os.path.join(target, FILES["detail"]),
                ["episode_id", "unit_code", "separation_date", "complication_sequence",
                 "complication_code", "onset_date"], out)


def write_series(target, expected, observed, inscope):
    """The reporting system's own quarterly series for the in-scope units: a rollup of the extract."""
    out = []
    for u in W.UNITS:
        if u[4] != "acute":
            continue
        for q in sim.quarters_for(u):
            label, fy, _, _, _ = W.QUARTERS[q]
            e = expected[(u[0], q)]
            o = observed[(u[0], q)]
            out.append([u[0], fy, label, inscope[(u[0], q)], o, round(e, 3), round(o / e, 4),
                        round(o / e * W.REF_RATE, 1)])
    out.sort(key=lambda r: (r[0], r[2]))
    return _csv(os.path.join(target, FILES["series"]),
                ["unit_code", "fiscal_year", "quarter", "in_scope_episodes", "observed_complications",
                 "expected_complications", "standardised_ratio", "adjusted_rate_per_1000"], out)


def write_summary(target, rows):
    """Separations and affected episodes by unit and month, as the activity reporting system publishes them:
    every care type, every unit, one row per unit and month, counted on distinct separations."""
    agg = {}
    for r in rows:
        key = (r["sep_date"].replace(day=1), r["unit"])
        cur = agg.setdefault(key, [0, 0])
        cur[0] += 1
        cur[1] += 1 if r["hac"] else 0
    out = []
    for (month, unit), (seps, hac) in sorted(agg.items()):
        out.append([month.isoformat(), unit, facility(unit, month), seps, hac])
    return _csv(os.path.join(target, FILES["summary"]),
                ["month", "unit_code", "facility", "separations", "episodes_with_complication"], out)


PLAN_NOTES = {
    "41210": "winter medical escalation beds stood down from 1 April",
    "41220": "two additional elective lists a week from 1 April",
    "41520": "elective general surgery list transferred in from Marrowfield",
    "41310": "ward refurbishment, four beds closed for the quarter",
}


def write_plan(target, plans):
    out = []
    for u in W.UNITS:
        if u[0] not in plans:
            continue
        out.append([u[0], W.EMBED_QUARTER, plans[u[0]], PLAN_NOTES.get(u[0], "")])
    out.sort()
    return _csv(os.path.join(target, FILES["plan"]),
                ["unit_code", "quarter", "planned_separations", "planning_note"], out)


DIVISION = {"gen_med": "Medicine", "resp": "Medicine", "renal": "Medicine", "aged": "Aged Care and Rehabilitation",
            "neuro": "Medicine", "subacute": "Aged Care and Rehabilitation", "gen_surg": "Surgery",
            "vasc": "Surgery", "colorectal": "Surgery", "ortho": "Surgery"}
SERVICE_LABEL = {"acute": "Acute inpatient", "rehabilitation": "Sub-acute rehabilitation",
                 "psychiatric": "Psychiatric inpatient"}


def write_units(target):
    """The unit master as it stands at the export date."""
    out = []
    for code, name, hosp, mix, service, _, _, _, _, _, status in W.UNITS:
        site = W.REALIGNED[code][1] if code in W.REALIGNED else hosp
        opened = "2019-07-01" if status != "merged_into" else "2016-07-01"
        closed = ""
        into = ""
        if status == "merged_into":
            closed = (W.MERGE_DATE - dt.timedelta(days=1)).isoformat()
            into = W.MERGE_INTO
            opened = "2014-07-01"
        if code == W.MERGE_INTO:
            opened = W.MERGE_DATE.isoformat()
        cis = "N" if code in W.DETAIL_EXCLUDED else "Y"
        out.append([code, name, site, SERVICE_LABEL[service], DIVISION[mix], cis, opened, closed, into])
    out.sort()
    return _csv(os.path.join(target, FILES["units"]),
                ["unit_code", "unit_name", "hospital_code", "service_type", "division",
                 "clinical_information_system", "opened_date", "closed_date", "amalgamated_into"], out)


def write_sites(target):
    beds = {"H1": 612, "H2": 388, "H3": 301, "H4": 112}
    return _csv(os.path.join(target, FILES["sites"]),
                ["hospital_code", "hospital_name", "peer_group", "overnight_beds"],
                [[c, n, p, beds[c]] for c, n, p in W.HOSPITALS])


def write_bands(target):
    return _csv(os.path.join(target, FILES["bands"]), ["drg_code", "drg_description", "risk_band"],
                [[c, d, b or ""] for c, d, b in W.DRGS])


def write_strata(target):
    out = [[b, a, "%.4f" % W.STRATUM_RATE[(b, a)]]
           for b in ["ORTH-MAJ", "ORTH-MIN", "SURG-MAJ", "SURG-MIN", "MED-COMPLEX", "MED-GEN"]
           for a in W.AGE_BANDS]
    return _csv(os.path.join(target, FILES["strata"]), ["risk_band", "age_band", "expected_rate"], out)


def write_catlist(target):
    out = [[c, n, g, W.GROUP_NAME[g]] for c, n, g in W.LIST_2026]
    return _csv(os.path.join(target, FILES["catlist"]),
                ["complication_code", "complication_name", "group_id", "group_name"], out)


def write_placements(target):
    """The improvement partnership's placement register: what the team did, cycle by cycle."""
    out = []
    cycle = 1
    rows = [("2023-24 Q4", "41360")] + [(W.QLABELS[k], c) for k, c in sorted(W.EMBEDS.items())] + \
           [(W.EMBED_CURRENT[0], W.EMBED_CURRENT[1])]
    for label, code in rows:
        fy, q = label.split(" Q")
        y0 = 2000 + int(fy.split("-")[0]) % 100
        start_month = {1: (7, 0), 2: (10, 0), 3: (1, 1), 4: (4, 1)}[int(q)]
        start = dt.date(y0 + start_month[1], start_month[0], 1)
        em, ey = (start.month + 3, start.year) if start.month <= 9 else (start.month - 9, start.year + 1)
        end = dt.date(ey, em, 1) - dt.timedelta(days=1)
        decided = start - dt.timedelta(days=start.day + 7)
        current = label == W.EMBED_CURRENT[0]
        out.append(["IP-%02d" % cycle, label, code, decided.isoformat(), start.isoformat(),
                    end.isoformat(), "%.1f" % W.TEAM_FTE,
                    "" if current else (end + dt.timedelta(days=21)).isoformat(),
                    "in progress" if current else "closed"])
        cycle += 1
    return _csv(os.path.join(target, FILES["placements"]),
                ["cycle", "assessment_quarter", "unit_code", "decision_date", "start_date", "end_date",
                 "team_fte", "close_out_date", "status"], out)


def write_dashboard(target, expected, observed, inscope, rows):
    """The quality dashboard's quarterly export: the position at the quarter ended 31 December 2026."""
    q = len(W.QUARTERS) - 1
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quarter position"
    head = ["Hospital", "Cost centre", "Unit", "In-scope episodes", "Complications",
            "Expected complications", "Standardised ratio", "Adjusted rate per 1,000",
            "Unplanned readmissions per 100", "Average length of stay (days)",
            "In-hospital mortality per 1,000"]
    ws.append(["Mid-Rivers Health Network: inpatient quality indicators"])
    ws.append(["Quarter ended 31 December 2026, extracted 5 March 2027. Quarter-end position; not a "
               "projection of any later quarter."])
    ws.append([])
    ws.append(head)
    sites = {c: n for c, n, _ in W.HOSPITALS}
    los = {}
    deaths = {}
    for r in rows:
        if r["q"] != q or not r["in_scope"]:
            continue
        cur = los.setdefault(r["unit"], [0, 0])
        cur[0] += r["los"]
        cur[1] += 1
    body = []
    for u in W.UNITS:
        if u[4] != "acute" or q not in sim.quarters_for(u):
            continue
        code = u[0]
        e, o, n = expected[(code, q)], observed[(code, q)], inscope[(code, q)]
        rr = sim.rng("dash", code)
        body.append([sites[facility(code, W.QUARTERS[q][2])], code, u[1], n, o, round(e, 1),
                     round(o / e, 3), round(o / e * W.REF_RATE, 1),
                     round(float(6.1 + 2.4 * rr.random()), 1),
                     round(los[code][0] / los[code][1], 1),
                     round(float(3.8 + 5.2 * rr.random()), 1)])
    body.sort(key=lambda r: (r[0], r[2]))
    for r in body:
        ws.append(r)
    bold = Font(bold=True)
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"].font = Font(italic=True, size=9)
    for c in range(1, len(head) + 1):
        cell = ws.cell(row=4, column=c)
        cell.font = bold
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.fill = PatternFill("solid", fgColor="E8E8E4")
    widths = [22, 11, 28, 11, 11, 12, 11, 12, 13, 12, 13]
    for i, wdt in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = "A5"
    for row in ws.iter_rows(min_row=5, min_col=4, max_col=11):
        for cell in row:
            cell.number_format = "0.0" if cell.column >= 6 else "#,##0"
    notes = wb.create_sheet("Notes")
    for line in [["Indicator", "Definition"],
                 ["Standardised ratio", "Complications observed divided by complications expected under "
                  "specification 4.0, in-scope episodes only"],
                 ["Adjusted rate per 1,000", "Standardised ratio times the reference rate of 96.5 "
                  "complications per 1,000 episodes"],
                 ["Unplanned readmissions per 100", "Unplanned readmissions within 28 days per 100 "
                  "separations, all care types"],
                 ["In-hospital mortality per 1,000", "Deaths in hospital per 1,000 separations, all care "
                  "types, not risk adjusted"],
                 ["Coverage", "Acute inpatient units only. Sub-acute and psychiatric units are outside the "
                  "complication indicator's scope."],
                 ["Owner", "Clinical Information and Coding, Mid-Rivers Health Network"],
                 ["Refreshed", "5 March 2027, for the quarter ended 31 December 2026"]]:
        notes.append(line)
    notes.column_dimensions["A"].width = 32
    notes.column_dimensions["B"].width = 86
    for c in ("A1", "B1"):
        notes[c].font = bold
    for row in notes.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    path = os.path.join(target, FILES["dashboard"])
    wb.save(path)
    return path


def write_distractors(target):
    """Three unit-level files the placement decision never reads."""
    out = []
    for u in W.UNITS:
        for q in sim.quarters_for(u)[-6:]:
            r = sim.rng("exp", u[0], q)
            out.append([u[0], W.QUARTERS[q][0], round(float(78 + 14 * r.random()), 1),
                        round(float(82 + 12 * r.random()), 1), int(28 + r.integers(0, 90))])
    _csv(os.path.join(target, FILES["experience"]),
         ["unit_code", "quarter", "overall_care_rating", "would_recommend_pct", "responses"], out)
    out = []
    for u in W.UNITS:
        for q in sim.quarters_for(u)[-4:]:
            for m in W.QUARTERS[q][4]:
                r = sim.rng("nurse", u[0], m.isoformat())
                out.append([m.isoformat(), u[0], round(float(4.1 + 1.9 * r.random()), 2),
                            round(float(0.03 + 0.14 * r.random()), 3),
                            round(float(0.02 + 0.07 * r.random()), 3)])
    _csv(os.path.join(target, FILES["nursing"]),
         ["month", "unit_code", "nursing_hours_per_patient_day", "agency_share", "overtime_share"], out)
    out = []
    for u in W.UNITS:
        for period in ("2026 Audit 1", "2026 Audit 2", "2026 Audit 3"):
            r = sim.rng("hh", u[0], period)
            moments = int(120 + r.integers(0, 180))
            out.append([u[0], period, moments, int(moments * float(0.78 + 0.18 * r.random())),
                        round(float(0.78 + 0.18 * r.random()) * 100, 1)])
    _csv(os.path.join(target, FILES["hygiene"]),
         ["unit_code", "audit_period", "moments_observed", "moments_compliant", "compliance_pct"], out)
    return [FILES["experience"], FILES["nursing"], FILES["hygiene"]]


DICTIONARY = [
    ("episode_outcomes_extract.csv", "episode_id", "Identifier of the episode record in the warehouse, "
     "allocated in one network-wide sequence by separation month", "code"),
    ("episode_outcomes_extract.csv", "unit_code", "Cost centre of the unit the episode separated from", "code"),
    ("episode_outcomes_extract.csv", "patient_uid", "Patient master index number", "code"),
    ("episode_outcomes_extract.csv", "admission_date", "Date the episode of care began", "date"),
    ("episode_outcomes_extract.csv", "separation_date", "Date the episode of care ended; an episode is "
     "reported in the period its separation falls in", "date"),
    ("episode_outcomes_extract.csv", "care_type", "Care type of the episode: Acute, Rehabilitation, "
     "Maintenance or Psychiatric", "code"),
    ("episode_outcomes_extract.csv", "los_days", "Nights stayed; zero for a same-day separation", "days"),
    ("episode_outcomes_extract.csv", "drg_code", "Diagnosis related group assigned by the grouper; 960Z and "
     "961Z are error groups", "code"),
    ("episode_outcomes_extract.csv", "age_band", "Age band at admission", "code"),
    ("episode_outcomes_extract.csv", "comorbidity_score", "Charlson comorbidity score where coding has "
     "keyed it; blank where it has not", "score"),
    ("episode_outcomes_extract.csv", "hac_recorded", "1 where the episode recorded one or more "
     "hospital-acquired complications", "indicator"),
    ("episode_outcomes_extract.csv", "hac_category_code", "Code of the first complication recorded against "
     "the episode, on the national list in force at separation", "code"),
    ("hac_detail_extract.csv", "complication_code", "Code of the complication on the national list in force "
     "at separation", "code"),
    ("hac_detail_extract.csv", "complication_sequence", "Order the complication was recorded in for the "
     "episode", "count"),
    ("hac_detail_extract.csv", "onset_date", "Date the complication was first documented", "date"),
    ("unit_indicator_series.csv", "standardised_ratio", "Complications observed divided by complications "
     "expected for the quarter, specification 4.0", "ratio"),
    ("unit_indicator_series.csv", "adjusted_rate_per_1000", "Standardised ratio times the reference rate",
     "complications per 1,000 episodes"),
    ("activity_summary_monthly.csv", "facility", "Hospital site the unit reported under in that month", "code"),
    ("activity_summary_monthly.csv", "separations", "Distinct separations the unit reported in the month, "
     "all care types", "count"),
    ("activity_summary_monthly.csv", "episodes_with_complication", "Separations in the month that recorded "
     "one or more complications", "count"),
    ("activity_plan_2027q4.csv", "planned_separations", "Separations planned for the unit for the quarter, "
     "all care types", "count"),
    ("unit_register.csv", "hospital_code", "Hospital site the unit is aligned to at the export date", "code"),
    ("unit_register.csv", "clinical_information_system", "Y where the unit's wards are on the clinical "
     "information system, which is the source of the complication level detail; N where a complication "
     "reaches the warehouse only as the episode record's own category", "indicator"),
    ("unit_register.csv", "service_type", "Service type of the unit", "code"),
    ("unit_register.csv", "amalgamated_into", "Cost centre the unit's activity transferred to on closure",
     "code"),
    ("hac_stratum_rates.csv", "expected_rate", "Complications expected per in-scope episode in the stratum, "
     "published with specification 4.0", "rate"),
    ("improvement_placements.csv", "assessment_quarter", "Quarter the team was embedded with the unit",
     "quarter"),
    ("improvement_placements.csv", "decision_date", "Date the placement was decided", "date"),
    ("improvement_placements.csv", "team_fte", "Full-time equivalent improvement staff embedded", "fte"),
]
GRAINS = [
    ("episode_outcomes_extract.csv", "Grain", "One row per episode record received from a source system. "
     "A separation is distinct within patient, admission date and diagnosis related group.", "-"),
    ("hac_detail_extract.csv", "Grain", "One row per complication recorded against an episode on the "
     "clinical information system.", "-"),
    ("unit_indicator_series.csv", "Grain", "One row per in-scope unit and closed quarter.", "-"),
    ("activity_summary_monthly.csv", "Grain", "One row per unit and month, all care types.", "-"),
    ("improvement_placements.csv", "Grain", "One row per placement cycle.", "-"),
]


def write_dictionary(target):
    return _csv(os.path.join(target, FILES["dictionary"]),
                ["file_name", "column_name", "description", "unit_of_measure"], DICTIONARY + GRAINS)


MANIFEST = [
    ("episode_outcomes_extract.csv", "Episode level extract of separations with their complication flags",
     "1 July 2024 to 31 December 2026, every unit", "episodes not yet coded; the quarter in progress",
     "Clinical Information and Coding"),
    ("hac_detail_extract.csv", "Complication level detail from the clinical information system",
     "the units on the clinical information system, 1 July 2024 to 31 December 2026",
     "units not on the clinical information system", "Clinical Information and Coding"),
    ("unit_indicator_series.csv", "Quarterly complication indicator series as reported",
     "in-scope units, specification 4.0 period", "quarters before 1 July 2024, reported on specification 3.2",
     "Clinical Governance"),
    ("activity_summary_monthly.csv", "Monthly activity and complication counts as published",
     "every unit and month in the period, all care types", "nothing", "Activity Reporting"),
    ("activity_plan_2027q4.csv", "Planned separations for the quarter from 1 April 2027",
     "every reporting unit", "sub-acute and psychiatric units are planned separately",
     "Performance and Planning"),
    ("unit_register.csv", "Unit master as it stands at the export date", "every cost centre in the period",
     "earlier site alignments and names", "Finance Systems"),
    ("hospital_sites.csv", "Hospital site reference", "the network's four hospitals", "nothing",
     "Finance Systems"),
    ("drg_risk_bands.csv", "Diagnosis related group to risk band map, specification 4.0",
     "every group in the period", "groups retired before the specification took effect", "Clinical Governance"),
    ("hac_stratum_rates.csv", "Published stratum rates, specification 4.0", "every stratum", "nothing",
     "Clinical Governance"),
    ("hac_category_list.csv", "National complication list, 2026 revision",
     "codes in force from 1 July 2026", "the 2022 revision, which the specification's appendix carries",
     "Clinical Governance"),
    ("improvement_placements.csv", "Improvement partnership placement register",
     "every placement since the partnership began", "the placement in progress has no close-out",
     "Clinical Improvement"),
    ("quality_dashboard_2026q2.xlsx", "Quality dashboard export, quarter ended 31 December 2026",
     "acute inpatient units", "sub-acute and psychiatric units", "Clinical Governance"),
    ("improvement_partnership_charter_v3.pdf", "Improvement partnership charter, version 3",
     "the partnership's purpose, placement and assessment clauses", "schedules 1 to 3",
     "Clinical Improvement"),
    ("hac_indicator_specification_v4.pdf", "Hospital-acquired complication indicator specification 4.0",
     "definition, risk adjustment, reporting and the complication list appendices", "worked examples",
     "Clinical Governance"),
    ("placement_cycle_notes.txt", "Notes on the placement cycle from the improvement unit's own thread",
     "the cycle from 1 April 2027", "nothing", "Clinical Improvement"),
    ("data_dictionary.csv", "Field semantics and grain statements for the extracts", "the files listed",
     "the reference files", "Clinical Information and Coding"),
    ("patient_experience_unit_quarterly.csv", "Patient experience survey results by unit",
     "the last six quarters", "units with fewer than 25 responses are reported as published",
     "Consumer Engagement"),
    ("nursing_hours_by_unit_monthly.csv", "Nursing hours per patient day and agency share by unit",
     "the last four quarters", "medical staffing", "Nursing and Midwifery"),
    ("hand_hygiene_audit_2026.csv", "Hand hygiene audit moments and compliance by unit",
     "the three 2026 audit periods", "audits before 2026", "Infection Prevention"),
]


def write_manifest(target):
    rows = [[f, c, cov, om, owner, W.PULL_DATE.isoformat(),
             "Mid-Rivers Health Network administrative and clinical systems", "Internal use"]
            for f, c, cov, om, owner in MANIFEST]
    rows.append([FILES["manifest"], "This inventory", "every file in the folder", "nothing",
                 "Clinical Information and Coding", W.PULL_DATE.isoformat(),
                 "Mid-Rivers Health Network administrative and clinical systems", "Internal use"])
    return _csv(os.path.join(target, FILES["manifest"]),
                ["file_name", "content", "covers", "omits", "prepared_by", "export_date", "source",
                 "licence"], rows)
