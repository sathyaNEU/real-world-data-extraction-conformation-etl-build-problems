"""The golden deliverables, computed from the shipped pack alone: the 2027 cycling threshold workbook (the setting,
the expected 2027 cycling days at each menu setting, the 2026 operation by month and the cycle chart) and the conformed
daily-spread file; and generator/figures.json, every figure the write-up states."""
import calendar
import csv
import datetime as dt
import io
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import methods as M    # noqa: E402
import world as W      # noqa: E402

MONTHS = [calendar.month_name[m] for m in M.OPS_MONTHS]


def compute(target):
    d = M.daily_spreads(target)
    g = M.golden(d)
    R, coef = M.rungs(d)
    cyc, lost = M.warranty_cycles(target)
    net = M.dam_net(target)
    ov = M.overlap(d)
    return dict(d=d, g=g, R=R, coef=coef, cyc=cyc, lost=lost, net=net, ov=ov)


def lost_spans(lost):
    """The dropped hours the meter shows the battery discharging in, as UTC spans."""
    spans = []
    for t in lost:
        if spans and t - spans[-1][1] == dt.timedelta(hours=1):
            spans[-1][1] = t
        else:
            spans.append([t, t])
    return spans


def write_csv(X, path):
    d = X["d"]
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(["operating_date", "settlement_point", "hours_in_day", "spread_usd_per_mwh", "set_at_hub_price"])
    rows = 0
    for r in d.itertuples():
        w.writerow([r.day.isoformat(), W.HUB, int(r.hours), "%.2f" % r.hub_spread, "N"])
        rows += 1
        if r.node_spread == r.node_spread:
            w.writerow([r.day.isoformat(), W.NODE, int(r.hours), "%.2f" % r.node_spread, "Y" if r.at_hub_price else "N"])
            rows += 1
    with open(path, "w", newline="", encoding="utf-8") as f:
        f.write(buf.getvalue())
    return rows


def write_workbook(X, path):
    import openpyxl
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    g, cyc, net = X["g"], X["cyc"], X["net"]
    call = g["call"]
    below = max(s for s in M.MENU if s < call)
    wb = openpyxl.Workbook()
    ink = Font(name="Calibri", size=11)
    bold = Font(name="Calibri", size=11, bold=True)
    title = Font(name="Calibri", size=14, bold=True, color="8A3B12")
    grey = Font(name="Calibri", size=9, italic=True, color="5A5F66")
    head = PatternFill("solid", fgColor="FBE5D6")
    pick = PatternFill("solid", fgColor="FFF2CC")
    thin = Border(bottom=Side(style="thin", color="9AA0A6"))

    ws = wb.active
    ws.title = "2027 threshold"
    ws["A1"] = "Redrock Draw 2027 cycling threshold: $%d per MWh" % call
    ws["A1"].font = title
    ws["A2"] = ("Expected 2027 cycling days at each setting on the trading system's menu (TRD-STOR-PARAM-07), against the "
                "warranty's Annual Cycle Limit of 300 Full Cycles; the lowest setting within it is the threshold (policy "
                "TSP-OPS-014, 4.2).")
    ws["A2"].font = grey
    ws.append([])
    hdr = ["Setting ($/MWh)", "Expected 2027 cycling days", "Within 300"]
    ws.append(hdr)
    for j in range(1, 4):
        c = ws.cell(row=4, column=j)
        c.font = bold
        c.fill = head
        c.border = thin
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for s in M.MENU:
        e = g["exp"][s]
        ws.append([s, round(e), "Yes" if e <= M.LIMIT else "No"])
        r = ws.max_row
        ws.cell(row=r, column=1).number_format = '"$"0'
        ws.cell(row=r, column=2).number_format = "0"
        ws.cell(row=r, column=3).alignment = Alignment(horizontal="center")
        if s == call:
            for j in range(1, 4):
                ws.cell(row=r, column=j).fill = pick
                ws.cell(row=r, column=j).font = bold
    ws.append([])
    ws.append(["Threshold for 2027", "$%d per MWh" % call])
    ws.append(["Expected 2027 cycling days at $%d" % call, round(g["exp"][call])])
    ws.append(["Next setting down", "$%d per MWh: %d expected cycling days, over the 300 limit" % (below,
                                                                                                   round(g["exp"][below]))])
    for r in range(ws.max_row - 2, ws.max_row + 1):
        ws.cell(row=r, column=1).font = bold
    ws.append([])
    ov = X["ov"]
    notes = [
        "Basis. The node's daily spread over MPI's day-ahead record from %s to %s (%s operating days): the node's own "
        "spread on the %d days MPI priced it at its own settlement point (%s to %s), and on every other day %.2f times "
        "the North Hub's spread, the ratio of the node's total spread to the hub's over those %d days. Expected cycling "
        "days are 365 times the share of the record's days whose spread is at or above the setting." % (
            X["d"]["day"].min().isoformat(), X["d"]["day"].max().isoformat(), format(len(X["d"]), ","), len(ov),
            ov["day"].min().isoformat(), ov["day"].max().isoformat(), g["k"], len(ov)),
        "The node's own days since MPI first priced it are one windy stretch: a third of the hub's 2026 days had a spread "
        "under $8, against 11 to 16 per cent in 2019 and 2021 to 2025 (31 per cent in 2020), while its average spread was "
        "ordinary. Set from the node's own days alone, the threshold would be $15 (%d expected days at $15)." % round(
            X["R"]["R0 the node's own record"]["exp"][15]),
        "Spread: the average of the four highest hourly day-ahead prices on the operating day less the average of the "
        "four lowest (policy 2.2). The node's first 13 days, priced at the hub's price, are left out of the ratio.",
    ]
    for t in notes:
        ws.append([t])
        c = ws.cell(row=ws.max_row, column=1)
        c.font = grey
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=3)
        ws.row_dimensions[ws.max_row].height = 62
    ws.append([])
    ws.append(["North Hub by year", "Average daily spread ($/MWh)", "Days with a spread under $8 (%)"])
    hr = ws.max_row
    for j in range(1, 4):
        c = ws.cell(row=hr, column=j)
        c.font = bold
        c.fill = head
        c.border = thin
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    yr = X["d"]["day"].map(lambda v: v.year)
    for y in range(2019, 2027):
        v = X["d"].loc[yr == y, "hub_spread"]
        ws.append([("%d" % y) + (" (to 2 November)" if y == 2026 else ""), round(float(v.mean()), 2),
                   round(100.0 * float((v < 8).mean()))])
        ws.cell(row=ws.max_row, column=2).number_format = "0.00"
    ws.append(["Source: %s; the windy stretches (autumn 2019 to autumn 2020, autumn 2025 onward) carry the flat days; "
               "2021 fell back to the ordinary share." % W.FILES["hub"]])
    ws.cell(row=ws.max_row, column=1).font = grey
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 14
    ws.freeze_panes = "A5"

    ws2 = wb.create_sheet("2026 operation")
    ws2["A1"] = "Redrock Draw 2026 operation at the $%d threshold: Full Cycles and day-ahead net by month" % W.SETTING_2026
    ws2["A1"].font = title
    ws2["A2"] = ("Full Cycles as the warranty counts them (Schedule 4): DC discharge recorded by the BMS for each hour, the "
                 "Revenue Meter's delivered energy over 0.975 for the hours the BMS export lost, the capacity tests left "
                 "out; contract months in Central prevailing time. Day-ahead net: the latest statement line for each hour "
                 "and charge type, MPI's sign reversed, the administrative fee once per operating day.")
    ws2["A2"].font = grey
    ws2["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws2.merge_cells("A2:E2")
    ws2.row_dimensions[2].height = 58
    ws2.append([])
    hdr = ["Month (2026)", "Full Cycles", "Cumulative Full Cycles", "Pace to 300 by 31 December", "Day-ahead net after MPI fees ($ thousands)"]
    ws2.append(hdr)
    for j in range(1, 6):
        c = ws2.cell(row=4, column=j)
        c.font = bold
        c.fill = head
        c.border = thin
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws2.row_dimensions[4].height = 32
    cum = 0.0
    for m, name in zip(M.OPS_MONTHS, MONTHS):
        cum += float(cyc[m])
        end = dt.date(2026, m, calendar.monthrange(2026, m)[1])
        pace = 300.0 * end.timetuple().tm_yday / 365.0
        ws2.append([name, round(float(cyc[m]), 2), round(cum, 2), round(pace, 1), round(float(net[m]) / 1000.0, 1)])
        r = ws2.max_row
        ws2.cell(row=r, column=2).number_format = "0.00"
        ws2.cell(row=r, column=3).number_format = "0.00"
        ws2.cell(row=r, column=4).number_format = "0.0"
        ws2.cell(row=r, column=5).number_format = "#,##0.0"
    ws2.append(["January to October", round(float(cyc.sum()), 2), None, None, round(float(net.sum()) / 1000.0, 1)])
    r = ws2.max_row
    for j in range(1, 6):
        ws2.cell(row=r, column=j).font = bold
    ws2.cell(row=r, column=2).number_format = "0.00"
    ws2.cell(row=r, column=5).number_format = "#,##0.0"
    ws2.append([])
    per = float(net.sum()) / float(cyc.sum())
    ws2.append(["Day-ahead net per Full Cycle, January to October, from the unrounded totals ($)", round(per)])
    ws2.cell(row=ws2.max_row, column=1).font = bold
    ws2.cell(row=ws2.max_row, column=2).number_format = '"$"#,##0'
    ws2.append([])
    spans = lost_spans(X["lost"])
    ws2.append(["Hours the meter recorded the battery discharging while the BMS export shows zero energy and its last "
                "state held (UTC hour beginning), counted at the meter under warranty 4.3:"])
    ws2.cell(row=ws2.max_row, column=1).font = grey
    for a, b in spans:
        ws2.append(["%s to %s UTC (%d hour%s)" % (a.strftime("%Y-%m-%d %H:00"), (b + dt.timedelta(hours=1)).strftime(
            "%Y-%m-%d %H:00"), (b - a).seconds // 3600 + 1 + (b - a).days * 24, "" if b == a else "s")])
        ws2.cell(row=ws2.max_row, column=1).font = grey
    ws2.append(["Capacity tests left out: %s (TEST) and %s (MAINT after the June firmware update), Cordell reports "
                "CT-RDRW-2026-01 and -02." % (W.TESTS[0].isoformat(), W.TESTS[1].isoformat())])
    ws2.cell(row=ws2.max_row, column=1).font = grey
    for col, wdt in zip("ABCDE", (46, 14, 16, 18, 22)):
        ws2.column_dimensions[col].width = wdt
    ws2.freeze_panes = "A5"

    ws3 = wb.create_sheet("Cycle chart")
    ws3["A1"] = "2026 Full Cycles to date against the warranty's 300"
    ws3["A1"].font = title
    bar = BarChart()
    bar.type = "col"
    bar.title = "Cumulative 2026 Full Cycles at each month's end, against the pace to 300 by 31 December"
    bar.y_axis.title = "Full Cycles"
    bar.x_axis.title = "Month end, 2026"
    data = Reference(ws2, min_col=3, min_row=4, max_row=4 + len(M.OPS_MONTHS))
    cats = Reference(ws2, min_col=1, min_row=5, max_row=4 + len(M.OPS_MONTHS))
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(cats)
    line = LineChart()
    line.add_data(Reference(ws2, min_col=4, min_row=4, max_row=4 + len(M.OPS_MONTHS)), titles_from_data=True)
    line.set_categories(cats)
    bar += line
    bar.height = 9
    bar.width = 20
    bar.y_axis.scaling.min = 0
    bar.y_axis.scaling.max = 300
    ws3.add_chart(bar, "A3")

    ws4 = wb.create_sheet("Sources")
    rows = [["File", "Used for"],
            [W.FILES["hub"], "the North Hub's daily spreads, 2019 to 2026"],
            [W.FILES["node"], "the node's daily spreads since MPI first priced it"],
            [W.FILES["points"], "the node's two names and its first days at the hub's price"],
            [W.FILES["policy"], "the spread, the cycling rule, the menu and the threshold rule"],
            [W.FILES["warranty"], "the Annual Cycle Limit and how Full Cycles are counted"],
            [W.FILES["bms"], "the DC discharge by hour"],
            [W.FILES["meter"], "the delivered energy for the hours the BMS export lost"],
            [W.FILES["tests"], "the capacity tests' discharge windows"],
            [W.FILES["settle"], "the day-ahead net"]]
    for r in rows:
        ws4.append(r)
    for j in (1, 2):
        ws4.cell(row=1, column=j).font = bold
        ws4.cell(row=1, column=j).fill = head
    ws4.column_dimensions["A"].width = 60
    ws4.column_dimensions["B"].width = 64
    ws4.append([])
    ws4.append(["Prepared by %s, %s, for the 2027 operating parameters due 1 December 2026." % (
        W.PEOPLE["analyst"], W.TITLES["analyst"])])
    ws4.cell(row=ws4.max_row, column=1).font = grey
    wb.save(path)


def figures(X, csv_rows):
    g, R, cyc, net = X["g"], X["R"], X["cyc"], X["net"]
    call = g["call"]
    below = max(s for s in M.MENU if s < call)
    d = X["d"]
    yr = d["day"].map(lambda x: x.year)
    hub = {int(y): round(float(d.loc[yr == y, "hub_spread"].mean()), 2) for y in range(2019, 2027)}
    flat = {int(y): round(100 * float((d.loc[yr == y, "hub_spread"] < 8).mean())) for y in range(2019, 2027)}
    ov = X["ov"]
    out = dict(
        call=call, below=below, limit=M.LIMIT, setting_2026=W.SETTING_2026,
        exp={str(s): round(g["exp"][s]) for s in M.MENU},
        exp_unrounded={str(s): g["exp"][s] for s in M.MENU},
        k=round(g["k"], 2), k_unrounded=g["k"], overlap_days=len(ov), record_days=len(d),
        node_days=int(d["node_spread"].notna().sum()), hub_priced_days=int(d["at_hub_price"].sum()),
        rungs={n: dict(call=r["call"], exp={str(s): round(v) for s, v in r["exp"].items()}) for n, r in R.items()},
        premium=round(float((ov["node_spread"] - ov["hub_spread"]).mean()), 2),
        hub_mean_by_year=hub, hub_flat_share_by_year=flat,
        cycles={str(m): round(float(cyc[m]), 2) for m in M.OPS_MONTHS},
        cycles_unrounded={str(m): float(cyc[m]) for m in M.OPS_MONTHS},
        cycles_total=round(float(cyc.sum()), 2),
        net_k={str(m): round(float(net[m]) / 1000.0, 1) for m in M.OPS_MONTHS},
        net_unrounded={str(m): float(net[m]) for m in M.OPS_MONTHS},
        net_total_k=round(float(net.sum()) / 1000.0, 1),
        per_cycle=int(round(float(net.sum()) / float(cyc.sum()))),
        lost_hours=len(X["lost"]), lost_spans=len(lost_spans(X["lost"])),
        csv_rows=csv_rows, tests=[t.isoformat() for t in W.TESTS], fee=0.275,
        cum_cycles={str(m): round(float(cyc.loc[:m].sum()), 2) for m in M.OPS_MONTHS},
        menu=list(M.MENU), meter_factor=M.TO_DC, repriced_days=[20, 21], contract_year=2027,
    )
    return out


def main(target, out_dir, fig_path):
    os.makedirs(out_dir, exist_ok=True)
    X = compute(target)
    write_workbook(X, os.path.join(out_dir, W.GOLDEN["workbook"]))
    n = write_csv(X, os.path.join(out_dir, W.GOLDEN["daily"]))
    F = figures(X, n)
    with open(fig_path, "w") as f:
        json.dump(F, f, indent=1, sort_keys=True)
        f.write("\n")
    return X, F
