"""The pack's documents and registers: Tallgrove's day-ahead dispatch policy, the excerpt of the supply and warranty
agreement, Cordell's 2026 capacity test reports, the data dictionary and the extract manifest."""
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import LETTER                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import inch                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import pack as P    # noqa: E402
import world as W   # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("LiberationSerif", "Regular"), ("LiberationSerif-Bold", "Bold"), ("LiberationSerif-Italic", "Italic"),
              ("LiberationSans", "Regular"), ("LiberationSans-Bold", "Bold")):
    try:
        pdfmetrics.getFont(nm)
    except KeyError:
        pdfmetrics.registerFont(TTFont(nm, FONT_DIR + nm.split("-")[0] + "-" + f + ".ttf"))
INK = colors.HexColor("#1d1f21")
GREY = colors.HexColor("#5a5f66")
RUST = colors.HexColor("#8a3b12")
NAVY = colors.HexColor("#1f3a5f")
BODY = ParagraphStyle("body", fontName="LiberationSerif", fontSize=10.5, leading=14.5, textColor=INK, spaceAfter=6)
NUM = ParagraphStyle("num", parent=BODY, leftIndent=30, firstLineIndent=-30)
H1 = ParagraphStyle("h1", fontName="LiberationSans-Bold", fontSize=14, leading=18, textColor=RUST, spaceAfter=2)
H1N = ParagraphStyle("h1n", parent=H1, textColor=NAVY)
H1S = ParagraphStyle("h1s", fontName="LiberationSans", fontSize=10, leading=13, textColor=GREY, spaceAfter=10)
H2 = ParagraphStyle("h2", fontName="LiberationSans-Bold", fontSize=10.5, leading=14, textColor=INK, spaceBefore=7, spaceAfter=3)
CELL = ParagraphStyle("cell", fontName="LiberationSans", fontSize=8.8, leading=11, textColor=INK)
CELLB = ParagraphStyle("cellb", parent=CELL, fontName="LiberationSans-Bold")
SMALL = ParagraphStyle("small", fontName="LiberationSans", fontSize=8.5, leading=11, textColor=GREY)


def _doc(path, title, author, footer):
    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFont("LiberationSans", 7.5)
        canvas.setFillColor(GREY)
        canvas.drawString(1.0 * inch, 0.6 * inch, footer)
        canvas.drawRightString(LETTER[0] - 1.0 * inch, 0.6 * inch, "Page %d" % doc.page)
        canvas.restoreState()
    d = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=1.0 * inch, rightMargin=1.0 * inch, topMargin=0.9 * inch,
                          bottomMargin=0.95 * inch, title=title, author=author, subject="", creator="", producer="")
    return d, frame


def _table(rows, widths):
    data = [[Paragraph(str(c), CELLB if i == 0 else CELL) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa0a6")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eaed")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 4),
                           ("RIGHTPADDING", (0, 0), (-1, -1), 4)]))
    return t


# ---------------------------------------------------------------------------------------------- the dispatch policy
def write_policy(path):
    doc, frame = _doc(path, "Redrock Draw Energy Storage: Day-Ahead Dispatch Policy", W.ORG,
                      "%s  |  TSP-OPS-014 Day-Ahead Dispatch Policy, Redrock Draw Energy Storage  |  Version 2.0, "
                      "approved 6 February 2026" % W.ORG)
    s = [Paragraph("Redrock Draw Energy Storage: Day-Ahead Dispatch Policy", H1),
         Paragraph("%s. Policy TSP-OPS-014, version 2.0, approved by the management team on 6 February 2026. Policy "
                   "owner: %s, %s." % (W.ORG, W.PEOPLE["asset"], W.TITLES["asset"]), H1S)]
    P_ = lambda n, t: s.append(Paragraph("<b>%s</b>&nbsp;&nbsp;%s" % (n, t), NUM))
    s.append(Paragraph("1. Scope", H2))
    P_("1.1", "This policy governs how %s (the Battery, 100 MW and 400 MWh, in commercial operation since 29 December "
              "2025) is cycled in the %s's day-ahead market, and how the cycling threshold is set for each contract "
              "year. %s schedules the Battery as its qualified scheduling entity (QSE) under the Trading Services "
              "Agreement." % (W.PROJECT, W.MARKET, W.QSE))
    s.append(Paragraph("2. Definitions", H2))
    P_("2.1", "<b>Operating day:</b> the market's operating day, midnight to midnight in Central prevailing time, its "
              "hours named by the hour in which they end (hour ending 1 to 24); it has 23 hours on the day clocks go "
              "forward and 25 on the day they go back.")
    P_("2.2", "<b>Spread:</b> for an operating day and a settlement point, the average of the four highest hourly "
              "day-ahead settlement point prices %s publishes for that operating day less the average of the four "
              "lowest, in dollars per MWh." % W.OPERATOR_SHORT)
    P_("2.3", "<b>The Battery's settlement point:</b> the resource node %s prices for the Battery (%s)." %
              (W.OPERATOR_SHORT, W.NODE_OLD))
    P_("2.4", "<b>Contract year:</b> a calendar year in Central prevailing time, as in the Supply and Warranty Agreement "
              "with %s." % W.OEM)
    s.append(Paragraph("3. Cycling", H2))
    P_("3.1", "The Battery cycles once in the day-ahead market on each operating day whose spread at its settlement "
              "point is at or above the threshold in force (a cycling day): it charges in the four lowest-priced hours "
              "that come before the four highest-priced hours and discharges in those four. The QSE submits the "
              "Battery's offers accordingly and does not offer its energy on any other operating day.")
    P_("3.2", "Each cycle charges the Battery to full and discharges its Rated Energy, 400 MWh at its DC terminals, "
              "once: one Full Cycle under the warranty.")
    P_("3.3", "The Battery is not offered on the day of a capacity test the Supplier requests or while it is on outage. "
              "When %s instructs the Battery in an emergency, the QSE follows the instruction." % W.OPERATOR_SHORT)
    s.append(Paragraph("4. The threshold", H2))
    P_("4.1", "The trading system accepts thresholds from $5 to $40 per MWh in steps of $5 (parameter "
              "TRD-STOR-PARAM-07).")
    P_("4.2", "By 1 December each year the Asset Manager sets the threshold for the next contract year at the lowest "
              "setting on the menu at which the Battery's expected number of cycling days in that contract year does "
              "not exceed the warranty's Annual Cycle Limit of 300 Full Cycles.")
    P_("4.3", "The threshold goes to the QSE with the year's other operating parameters, which the Trading Services "
              "Agreement requires by 1 December for the contract year starting on 1 January.")
    P_("4.4", "For contract year 2026, the Battery's first, the threshold is $35 per MWh.")
    s.append(Paragraph("5. Cycles outside the threshold", H2))
    P_("5.1", "Capacity tests the Supplier requests and discharges %s instructs are not planned. The Asset Manager "
              "tracks the Full Cycles the warranty counts each contract month, whatever their cause, against the "
              "Annual Cycle Limit." % W.OPERATOR_SHORT)
    s.append(Paragraph("6. Records", H2))
    P_("6.1", "The operating records for this policy are %s's day-ahead price archive, the BMS hourly export from "
              "Cordell's portal, the revenue meter data %s posts for the Battery's meter and the QSE's settlement "
              "statement extracts." % (W.OPERATOR_SHORT, W.OPERATOR_SHORT))
    s.append(Spacer(1, 10))
    s.append(Paragraph("Supersedes version 1.0 (commissioning, 14 November 2025). Next review: with the 2027 parameters.",
                       SMALL))
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ----------------------------------------------------------------------------------- the warranty agreement excerpt
def write_warranty(path):
    doc, frame = _doc(path, "Supply and Warranty Agreement: excerpt", W.OEM,
                      "Supply and Warranty Agreement between %s and %s, dated 27 June 2025  |  Excerpt" % (W.OEM, W.ORG))
    s = [Paragraph("Supply and Warranty Agreement: Excerpt", H1N),
         Paragraph("Between %s (the Supplier) and %s (the Owner) for the battery energy storage system at %s, dated 27 "
                   "June 2025. Excerpt: Article 1 (selected definitions) and Schedule 4 (Performance Warranty)." %
                   (W.OEM, W.ORG, W.PROJECT), H1S)]
    P_ = lambda n, t: s.append(Paragraph("<b>%s</b>&nbsp;&nbsp;%s" % (n, t), NUM))
    s.append(Paragraph("Article 1. Definitions (selected)", H2))
    P_("1.4", "<b>BMS</b> means the battery management system the Supplier provides with the Battery, including its "
              "data portal.")
    P_("1.9", "<b>Capacity Test</b> means a test of the Battery's usable energy carried out at the Supplier's request "
              "under Schedule 5.")
    P_("1.11", "<b>Contract Year</b> means each calendar year beginning on the 1 January after the Commercial Operation "
               "Date, in the Project's local prevailing time; <b>Contract Month</b> means a calendar month within it.")
    P_("1.17", "<b>Full Cycle</b> means a quantity of energy discharged at the Battery's DC terminals equal to the Rated "
               "Energy; smaller discharges accumulate toward Full Cycles.")
    P_("1.23", "<b>Rated Energy</b> means 400 MWh discharged at the Battery's DC terminals.")
    P_("1.24", "<b>Revenue Meter</b> means the Project's settlement meter at the point of interconnection.")
    s.append(Paragraph("Schedule 4. Performance Warranty", H2))
    P_("4.1", "<b>Annual Cycle Limit.</b> The Performance Warranty is given on the basis that the Battery performs no "
              "more than 300 Full Cycles in any Contract Year.")
    P_("4.2", "<b>Counting.</b> Full Cycles are counted from the discharge energy at the DC terminals that the BMS "
              "records for each hour. Energy discharged during a Capacity Test is not counted.")
    P_("4.3", "<b>Data.</b> For any hour for which the BMS record is unavailable, the discharge energy for that hour is "
              "the energy delivered at the Revenue Meter in that hour divided by 0.975.")
    P_("4.4", "<b>Reporting.</b> Within fifteen days of the end of each calendar quarter the Owner reports to the "
              "Supplier the Full Cycles counted in each Contract Month of the quarter.")
    P_("4.5", "<b>Cycles above the limit.</b> Each Full Cycle above the Annual Cycle Limit in a Contract Year reduces "
              "the Guaranteed Energy Capacity for every later Contract Year by 0.035 per cent and is charged at the rate "
              "in the Long-Term Service Agreement.")
    P_("4.6", "<b>Capacity.</b> The usable energy measured in a Capacity Test shall be at least the Guaranteed Energy "
              "Capacity for the Contract Year (Contract Year 1: 400 MWh; reducing by 2.1 per cent a year).")
    s.append(Spacer(1, 10))
    s.append(Paragraph("Excerpt provided to the Owner's asset management team. Schedules 1 to 3 and 5 to 9 omitted.", SMALL))
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ------------------------------------------------------------------------------------------- the capacity test reports
def write_tests(path, dayrec):
    import datetime as dt
    doc, frame = _doc(path, "Capacity Test Reports 2026: Redrock Draw Energy Storage", W.OEM,
                      "%s  |  Field Services  |  Capacity test reports, Redrock Draw Energy Storage, 2026" % W.OEM)
    s = [Paragraph("Capacity Test Reports, 2026", H1N),
         Paragraph("%s, Field Services. Site: %s (BMS site RDRW-BESS-01). Procedure CES-CT-04 revision C. Times are "
                   "Central prevailing time." % (W.OEM, W.PROJECT), H1S)]
    for i, d in enumerate(W.TESTS, 1):
        r = dayrec[d]
        m = r["measured"]
        dis_min = int(round(m / 100.0 * 60))
        end = dt.datetime(d.year, d.month, d.day, 12) + dt.timedelta(minutes=dis_min)
        s.append(Paragraph("Report CT-RDRW-2026-%02d" % i, H2))
        rows = [["Item", "Result"],
                ["Requested by", "%s under Schedule 5 of the Supply and Warranty Agreement" % W.OEM],
                ["Charge", "%s 06:00 to 11:00, constant power to full" % d.isoformat()],
                ["Rest", "11:00 to 12:00"],
                ["Discharge", "%s 12:00 to %s, 100 MW at the DC terminals to empty" % (d.isoformat(), end.strftime("%H:%M"))],
                ["DC energy discharged", "%.1f MWh" % m],
                ["Measured usable energy", "%.1f MWh" % m],
                ["Guaranteed Energy Capacity, Contract Year 1", "400.0 MWh"],
                ["Result", "Pass"]]
        s.append(_table(rows, [2.4 * inch, 4.0 * inch]))
        s.append(Paragraph("Discharge: %s 12:00 to %s. Witnessed for the Owner by Tallgrove site operations. The "
                           "Battery was not offered to the market on the test day." % (d.isoformat(), end.strftime("%H:%M")),
                           BODY))
        s.append(Spacer(1, 6))
    s.append(Paragraph("The BMS takes the measured usable energy as its usable-energy setting after each test.", SMALL))
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ------------------------------------------------------------------------------------------------- the dictionary
def write_dictionary(path, counts):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    F = W.FILES
    c = lambda k: format(counts[k], ",")
    rows = [
        (F["hub"], "", "%s day-ahead settlement point prices for %s, from %s's yearly archive files for 2019 to 2026 "
                       "as downloaded and appended; one row per operating day and hour. %s rows." %
         (W.OPERATOR_SHORT, W.HUB, W.OPERATOR_SHORT, c("hub"))),
        (F["hub"], "DeliveryDate", "The operating day. The archive files to 2021 write it MM/DD/YYYY; from 2022, YYYY-MM-DD."),
        (F["hub"], "HourEnding", "The hour by the time it ends, Central prevailing time. To 2021 written 01:00 to 24:00; "
                                 "from 2022 written 1 to 24."),
        (F["hub"], "RepeatedHourFlag", "Marks the second hour ending 2 on the day clocks go back (Y or N to 2021, true "
                                       "or false from 2022)."),
        (F["hub"], "SettlementPointPrice", "Day-ahead settlement point price, dollars per MWh."),
        (F["node"], "", "%s day-ahead settlement point prices for the Battery's resource node from the first operating "
                        "day %s priced it; the same layout as the archive from 2022. %s rows." %
         (W.OPERATOR_SHORT, W.OPERATOR_SHORT, c("node"))),
        (F["points"], "", "%s's settlement point list as of the pull: hubs, load zones and resource nodes, with the "
                          "operating days each was priced and any predecessor." % W.OPERATOR_SHORT),
        (F["rt"], "", "%s real-time settlement point prices for %s, operating days 1 January to 1 November 2026, one row "
                     "per fifteen-minute interval. %s rows." %
         (W.OPERATOR_SHORT, W.HUB, c("rt"))),
        (F["as_prices"], "", "%s day-ahead ancillary service clearing prices (MCPC, dollars per MW per hour) for 2026: "
                             "regulation up and down, responsive reserve and non-spinning reserve. %s rows." %
         (W.OPERATOR_SHORT, c("as_prices"))),
        (F["bms"], "", "%s portal export of the Battery's BMS for UTC 1 January to 2 November 2026, one row per hour. %s "
                       "rows." % (W.OEM, c("bms"))),
        (F["bms"], "interval_start_utc", "Start of the hour, UTC."),
        (F["bms"], "operating_mode", "The BMS's mode in the hour: IDLE, CHARGE, DISCHARGE, TEST, MAINT (maintenance) or "
                                     "FAULT."),
        (F["bms"], "dc_charge_kwh, dc_discharge_kwh", "Energy into and out of the battery at its DC terminals in the hour, kWh."),
        (F["bms"], "pcs_ac_charge_kwh, pcs_ac_discharge_kwh", "The same energy at the power conversion system's AC "
                                                              "terminals, kWh."),
        (F["bms"], "soc_pct", "State of charge at the end of the hour, per cent of the usable energy the BMS holds for "
                              "the battery."),
        (F["bms"], "avg_cell_temperature_c", "Average cell temperature in the hour, degrees Celsius."),
        (F["meter"], "", "%s revenue meter data for the Battery's settlement meter at the point of interconnection, "
                         "as posted to the market participant portal for operating days to 1 November 2026; one row per "
                         "fifteen-minute interval. %s rows." %
         (W.OPERATOR_SHORT, c("meter"))),
        (F["meter"], "interval_ending_utc", "End of the fifteen-minute interval, UTC."),
        (F["meter"], "mwh_delivered", "Energy delivered to the grid in the interval, MWh."),
        (F["meter"], "mwh_received", "Energy taken from the grid in the interval (charging and the station's own load), "
                                     "MWh."),
        (F["settle"], "", "The QSE's extract of %s day-ahead settlement statement lines for the Battery's resource, "
                          "statements issued 1 January to 4 November 2026; one row per statement, hour and charge type. "
                          "%s rows." % (W.OPERATOR_SHORT, c("settle"))),
        (F["settle"], "statement_type", "INITIAL is issued two days after the operating day; FINAL thirty-one days "
                                        "after it, restating the operating day."),
        (F["settle"], "settlement_point", "The resource's settlement point as named when the statement was issued."),
        (F["settle"], "charge_type", "DAES day-ahead energy sale; DAEP day-ahead energy purchase; DAAF day-ahead "
                                     "administrative fee, $0.275 per MWh of the day-ahead energy awarded to the "
                                     "resource on the operating day."),
        (F["settle"], "bill_determinant_mwh", "The quantity the line's charge is computed on, MWh."),
        (F["settle"], "amount_usd", "Dollars, %s's sign: positive is owed by the QSE, negative is owed to it." %
         W.OPERATOR_SHORT),
        (F["policy"], "", "Tallgrove's day-ahead dispatch policy for the Battery, version 2.0."),
        (F["warranty"], "", "Excerpt of the Supply and Warranty Agreement with %s." % W.OEM),
        (F["tests"], "", "%s's reports of the 2026 capacity tests." % W.OEM),
    ]
    units = {"kwh": "kWh", "mwh": "MWh", "usd": "USD", "pct": "per cent", "price": "$/MWh", "mcpc": "$/MW-h",
             "temperature": "deg C"}

    def unit(field):
        f = field.lower()
        for k, u in units.items():
            if k in f:
                return u
        return None
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Field definitions"
    for j, h in enumerate(["Source file", "Column", "Meaning", "Unit"], 1):
        cc = ws.cell(row=1, column=j, value=h)
        cc.font = Font(bold=True)
        cc.fill = PatternFill("solid", fgColor="FBE5D6")
    for i, r in enumerate(rows, 2):
        for j, v in enumerate(list(r) + [unit(r[1]) if r[1] else None], 1):
            cc = ws.cell(row=i, column=j, value=v if v != "" else None)
            cc.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 92
    ws.column_dimensions["D"].width = 11
    ws.freeze_panes = "B2"
    wb.save(path)


def write_manifest(path, counts):
    F = W.FILES
    rows = [
        [F["hub"], "%s day-ahead settlement point prices, %s, yearly archive files 2019 to 2026" % (W.OPERATOR_SHORT, W.HUB),
         W.ARCHIVE_PULL.isoformat(), counts["hub"], "csv", "%s public market data" % W.OPERATOR_SHORT],
        [F["node"], "%s day-ahead settlement point prices, %s and %s" % (W.OPERATOR_SHORT, W.NODE_OLD, W.NODE),
         W.ARCHIVE_PULL.isoformat(), counts["node"], "csv", "%s public market data" % W.OPERATOR_SHORT],
        [F["points"], "%s settlement point list" % W.OPERATOR_SHORT, W.ARCHIVE_PULL.isoformat(), counts["points"], "csv",
         "%s public market data" % W.OPERATOR_SHORT],
        [F["rt"], "%s real-time settlement point prices, %s, 2026" % (W.OPERATOR_SHORT, W.HUB), W.ARCHIVE_PULL.isoformat(),
         counts["rt"], "csv", "%s public market data" % W.OPERATOR_SHORT],
        [F["as_prices"], "%s day-ahead ancillary service clearing prices, 2026" % W.OPERATOR_SHORT,
         W.ARCHIVE_PULL.isoformat(), counts["as_prices"], "csv", "%s public market data" % W.OPERATOR_SHORT],
        [F["bms"], "%s portal, BMS hourly export, site RDRW-BESS-01" % W.OEM, W.BMS_PULL.isoformat(), counts["bms"], "csv",
         "Owner's data under the Supply and Warranty Agreement"],
        [F["meter"], "%s market participant portal, revenue meter data, meter RDRW-POI-M1" % W.OPERATOR_SHORT,
         W.METER_PULL.isoformat(), counts["meter"], "csv", "market participant data"],
        [F["settle"], "%s, day-ahead settlement statement lines for resource %s" % (W.QSE, W.RESOURCE),
         W.SETTLE_PULL.isoformat(), counts["settle"], "csv", "Tallgrove confidential"],
        [F["policy"], "Tallgrove, TSP-OPS-014 Day-Ahead Dispatch Policy, version 2.0", W.POLICY_DATE.isoformat(), "", "pdf",
         "Tallgrove internal"],
        [F["warranty"], "Supply and Warranty Agreement with %s, excerpt" % W.OEM, W.WARRANTY_DATE.isoformat(), "", "pdf",
         "Tallgrove confidential"],
        [F["tests"], "%s Field Services, capacity test reports CT-RDRW-2026-01 and -02" % W.OEM, "2026-09-18", "", "pdf",
         "Tallgrove confidential"],
        [F["dict"], "Tallgrove asset management, field definitions for this folder", W.PACK_DATE.isoformat(), "", "xlsx",
         "Tallgrove internal"],
    ]
    out = [[r[0], r[1], r[2], r[3], r[5]] for r in rows]
    return P.write_csv(path, ["file", "description", "pulled_or_issued", "record_count", "access"], out)
