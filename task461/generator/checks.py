"""task461 checks: the three rules executed (R1 to R3c), every semantic fork (SF1 to SF12), the ladder, the asks'
devices and stops, the separation line, the distractors, the input gates and the realism sweep.

    python3 task461/generator/checks.py      (run by build.py after the build)
"""
import calendar
import datetime as dt
import json
import os
import re
import shutil
import sys
import tempfile
import zipfile

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import market as MK     # noqa: E402
import methods as M     # noqa: E402
import ops as O         # noqa: E402
import world as W       # noqa: E402

TASK = os.path.dirname(HERE)
RESULTS = []


def ok(label, cond, detail=""):
    RESULTS.append((label, bool(cond), detail))


def edge(v, unit):
    """Distance from the rounding edge at the given unit (0.01, 0.1, 1)."""
    x = v / unit
    return abs(x - np.floor(x) - 0.5) * unit


# ------------------------------------------------------------------------------------------------- price variants
def spread_table(target, mode=None):
    """The conformed daily record with one Layer 1 item mishandled."""
    pts = pd.read_csv(M.find(target, "mpi_settlement_point_list"), dtype=str).fillna("")
    succ = {r.predecessor: r.settlement_point for r in pts.itertuples() if r.predecessor}
    frames = []
    for stem in ("mpi_dam_settlement_point_prices_north_hub", "mpi_dam_settlement_point_prices_redrock"):
        p = M.read_prices(M.find(target, stem))
        if mode == "he24_next":
            p.loc[p["he"] == 24, "day"] = p.loc[p["he"] == 24, "day"].map(lambda d: d + dt.timedelta(days=1))
        if mode == "drop_repeated":
            p = p[~p["rep"]]
        if mode == "fill_missing":
            add = []
            for d, g in p.groupby(["day", "point"]):
                if len(g) == 23:
                    add.append(dict(day=d[0], he=3, rep=False, point=d[1], price=float(g.loc[g["he"].isin([2, 4]), "price"].mean())))
            p = pd.concat([p, pd.DataFrame(add)])
        if mode == "new_name_only" and "redrock" in stem:
            p = p[p["point"] == W.NODE]
        if mode != "new_name_only":
            p["point"] = p["point"].replace(succ)
        frames.append(p)
    hub, node = frames
    hub = hub[hub["day"] <= M.LAST_DAY]
    node = node[node["day"] <= M.LAST_DAY]
    hs = hub.groupby("day")["price"].apply(M.tb4).rename("hub_spread")
    ns = node.groupby("day")["price"].apply(M.tb4).rename("node_spread")
    d = pd.concat([hs, ns], axis=1).reset_index().rename(columns={"index": "day"})
    d = d[d["day"] >= dt.date(2019, 1, 1)]
    d["at_hub_price"] = d["day"].between(W.NODE_CREATED, W.NODE_PROVISIONAL_END - dt.timedelta(days=1)) & d[
        "node_spread"].notna()
    if mode == "hub_priced_kept":
        d["at_hub_price"] = False
    return d.reset_index(drop=True)


def call_and_sides(exp):
    return M.pick(exp), exp[15] > M.LIMIT, exp[20] <= M.LIMIT


def run(target, golden, x=None, meta_path=None, fig_path=None):
    F = json.load(open(fig_path or os.path.join(HERE, "figures.json")))
    meta = json.load(open(meta_path or os.path.join(TASK, "metadata.json")))
    files = sorted(os.listdir(target))
    d = M.daily_spreads(target)
    g = M.golden(d)
    R, coef = M.rungs(d)
    ov = M.overlap(d)
    H = d["hub_spread"].to_numpy()
    yr = d["day"].map(lambda v: v.year)

    # ------------------------------------------------------------------ the three rules (dataset-generation 7.1)
    heads = {f: open(os.path.join(target, f), encoding="utf-8").readline().lower() for f in files if f.endswith(".csv")}
    ok("R1 the golden reads the hourly price rows of both points",
       len(d) == 2863 and d["node_spread"].notna().sum() == 336, "%d hub days, %d node days" % (
           len(d), d["node_spread"].notna().sum()))
    ok("R1 no shipped column holds a daily spread, a carried spread or an expected count",
       not any(re.search(r"spread|expected|cycling_days|carried", h) for h in heads.values()))
    x = x or __import__("build").world()
    clean = x["days"]
    cn = {v["day"]: MK.tb4(v["node_pub"]) for v in clean if v["day"] >= W.NODE_CREATED}
    ch = {v["day"]: MK.tb4(v["hub"]) for v in clean}
    r0_clean = M.pick(M.expected_days(list(cn.values())))
    ok("R3a the default path on the clean truth (the node's own days, every filed rule applied) files $15",
       r0_clean == 15, "$%s" % r0_clean)
    ovc = [k for k in cn if k >= W.NODE_PROVISIONAL_END]
    kc = sum(cn[k] for k in ovc) / sum(ch[k] for k in ovc)
    days_c = sorted(ch)
    series_c = [cn[k] if k in cn and k >= W.NODE_PROVISIONAL_END else kc * ch[k] for k in days_c]
    ok("R3b the default path with only the decisive step changed (the extension through the overlap) files $20",
       M.pick(M.expected_days(series_c)) == 20, "$%s" % M.pick(M.expected_days(series_c)))
    prem_c = np.mean([cn[k] - ch[k] for k in ovc])
    ok("R3b each lower correction alone stays wrong: the hub's record $10, the hub plus the premium $30",
       M.pick(M.expected_days([ch[k] for k in days_c])) == 10 and
       M.pick(M.expected_days([ch[k] + prem_c for k in days_c])) == 30)
    # R3c: no copied or flat block; the only node days equal to the hub's are the 13 the list's note names
    hub_p = M.read_prices(M.find(target, "mpi_dam_settlement_point_prices_north_hub"))
    prof = hub_p.groupby("day")["price"].apply(lambda s: tuple(s.round(2)))
    ok("R3c no hub day repeats another day's prices and none is flat",
       prof.nunique() == len(prof) and hub_p.groupby("day")["price"].std().min() > 0.5)
    ok("R3c the only node days at the hub's prices are the 13 the settlement point list names", int(
        d["at_hub_price"].sum()) == 13 and set(d.loc[d["at_hub_price"], "day"]) == {
        W.NODE_CREATED + dt.timedelta(days=i) for i in range(13)})
    from pypdf import PdfReader
    doc_text = {f: re.sub(r"\s+", " ", " ".join(p.extract_text() for p in PdfReader(os.path.join(target, f)).pages))
                for f in files if f.endswith(".pdf")}
    pol = doc_text[W.FILES["policy"]]
    ok("R3c no document names a 2027 setting, an expected count or a basis for the expectation", not re.search(
        r"\$(10|15|20|25|30) per MWh|expected cycling days|ordinary|weather|North Hub's record|carried to|carried over",
        " ".join(doc_text.values())), "policy names only the menu and the 2026 threshold")
    ok("R3c the policy files the spread, the rule and the estimand", "four highest hourly day-ahead" in pol and
       "expected number of cycling days" in pol and "$35 per MWh" in pol)

    # ------------------------------------------------------------------------------------------- the ladder
    ok("golden: $20, $15's expected days over 300 and $20's within", g["call"] == 20 and g["exp"][15] > 300 and
       g["exp"][20] <= 300, "%.1f / %.1f" % (g["exp"][15], g["exp"][20]))
    lvl = lambda v: float(np.quantile(np.asarray(v, float), 1 - 300 / 365))
    gl = lvl(g["series"])
    ok("golden's 300-day level sits inside $15 to $20 with $1.5 or more to each edge", 16.5 <= gl <= 18.5, "%.2f" % gl)
    for name in ("R0 the node's own record", "R0 the node's record, the hub-priced days dropped", "R0 the node's 2026",
                 "Rx the node's record scaled by the hub's long-run mean"):
        ok("rung %s files $15" % name, R[name]["call"] == 15, "$%s" % R[name]["call"])
    ok("R0's 300-day level at $14 or under", lvl(d.loc[d["node_spread"].notna(), "node_spread"]) <= 14.0)
    ok("R1 (the hub plus the premium) files $30", R["R1 the hub plus the node's average spread premium"]["call"] == 30)
    ok("the hub's record alone files $10", R["Rx the hub's record alone"]["call"] == 10)
    # the hourly basis, by hour and by month and hour, added to the hub's hourly prices
    hubp = M.read_prices(M.find(target, "mpi_dam_settlement_point_prices_north_hub"))
    nodep = M.read_prices(M.find(target, "mpi_dam_settlement_point_prices_redrock"))
    j = nodep.merge(hubp, on=["day", "he", "rep"], suffixes=("_n", "_h"))
    j = j[j["day"].isin(set(ov["day"]))]
    j["basis"] = j["price_n"] - j["price_h"]
    j["month"] = j["day"].map(lambda v: v.month)
    bh = j.groupby("he")["basis"].mean()
    bmh = j.groupby(["month", "he"])["basis"].mean()
    hubp["month"] = hubp["day"].map(lambda v: v.month)
    hb = hubp.assign(p=hubp["price"] + hubp["he"].map(bh)).groupby("day")["p"].apply(M.tb4)
    hbm = hubp.assign(p=hubp["price"] + [bmh.get((m, h), bh[h]) for m, h in zip(hubp["month"], hubp["he"])]).groupby(
        "day")["p"].apply(M.tb4)
    ok("the hourly basis by hour files $30 and by month and hour $25 or more",
       M.pick(M.expected_days(hb)) == 30 and M.pick(M.expected_days(hbm)) >= 25,
       "$%s, $%s" % (M.pick(M.expected_days(hb)), M.pick(M.expected_days(hbm))))
    xs, ys = j["price_h"].to_numpy(), j["price_n"].to_numpy()
    slope = ((xs - xs.mean()) * (ys - ys.mean())).sum() / ((xs - xs.mean()) ** 2).sum()
    ok("an hourly price regression's slope on the hub's spreads files $15", M.pick(M.expected_days(slope * H)) == 15,
       "slope %.2f" % slope)
    ok("the extension's ratio is 2.21 and the overlap 323 days", round(g["k"], 2) == 2.21 and len(ov) == 323)

    # ---------------------------------------------------------------------------------------- the forks
    ok("SF1 the policy defines the spread (four highest less four lowest)", "average of the four highest hourly" in pol)
    ok("SF2 the price files carry no UTC clock", all("utc" not in heads[f] for f in heads if f.startswith("mpi_dam_settlement")))
    variants = {}
    x_ = d.copy()
    s_full = g["series"][yr.le(2025).to_numpy()]
    variants["SF3 full years 2019 to 2025"] = M.expected_days(s_full)
    per = {s: np.mean([365.0 * (g["series"][(yr == y).to_numpy()] >= s).mean() for y in range(2019, 2027)])
           for s in M.MENU}
    variants["SF3 each year's count weighted equally"] = per
    med = {s: float(np.median([365.0 * (g["series"][(yr == y).to_numpy()] >= s).mean() for y in range(2019, 2027)]))
           for s in M.MENU}
    variants["SF3 the median year"] = med
    for label, mode in (("SF4 the new name only", "new_name_only"), ("SF5 the hub-priced days kept", "hub_priced_kept"),
                        ("SF6 hour ending 24 read as the next day's", "he24_next"),
                        ("SF7 the repeated hour dropped", "drop_repeated"), ("SF7 the missing hour filled", "fill_missing")):
        dv = spread_table(target, mode)
        if mode == "new_name_only":
            dv.loc[dv["day"] < W.NODE_RENAME, "node_spread"] = np.nan
        variants[label] = M.golden(dv)["exp"]
    for label, exp in variants.items():
        c, a, b = call_and_sides(exp)
        ok("%s: same call, $15 over and $20 within" % label, c == 20 and a and b, "%.1f / %.1f" % (exp[15], exp[20]))
    ok("SF8 no setting's expected days within a day of 300", min(abs(v - 300) for v in g["exp"].values()) >= 1.0)
    # SF9: 2026's warranty cycles per day-ahead cycling day
    dr = x["dayrec"]
    cyc_days = [r for dd, r in dr.items() if r["status"] == "cycle" and W.OPS_FIRST <= dd <= W.OPS_LAST]
    per_day = np.mean([r["dc"] / 400.0 for r in cyc_days])
    ok("SF9 a day-ahead cycling day is one Full Cycle (1.00 to two decimals)", round(per_day, 2) == 1.00, "%.4f" % per_day)
    ests = {}
    x0, y0 = ov["hub_spread"].to_numpy(), ov["node_spread"].to_numpy()
    b0 = (x0 * y0).sum() / (x0 * x0).sum()
    ests["through the origin"] = b0 * H
    ests["regression with intercept"] = coef["a"] + coef["b"] * H
    vr = y0.std() / x0.std()
    ests["variance ratio"] = (y0.mean() - vr * x0.mean()) + vr * H
    qs = np.linspace(0.01, 0.99, 99)
    ests["quantile mapping"] = np.interp(H, np.quantile(x0, qs), np.quantile(y0, qs))
    ov_r = ov[ov["day"] >= W.NODE_RENAME]
    ests["the overlap from the rename"] = (ov_r["node_spread"].sum() / ov_r["hub_spread"].sum()) * H
    allnode = d[d["node_spread"].notna()]
    ests["the hub-priced days in the overlap"] = (allnode["node_spread"].sum() / allnode["hub_spread"].sum()) * H
    ests["2023 onward"] = (g["k"] * H)[yr.ge(2023).to_numpy()]
    ests["every day carried (no splice)"] = g["k"] * H
    for label, v in ests.items():
        e = M.expected_days(v)
        c, a, b = call_and_sides(e)
        ok("SF10 %s: same call" % label, c == 20 and a and b, "%.1f / %.1f" % (e[15], e[20]))
    ok("SF11 the call is the lowest setting within 300", g["call"] == min(s for s in M.MENU if g["exp"][s] <= 300))
    ok("SF12 contract year 2027 has 365 days", (dt.date(2028, 1, 1) - dt.date(2027, 1, 1)).days == 365)
    # the data-native structure
    fl = d.groupby(yr)["hub_spread"].apply(lambda v: (v < 8).mean())
    mn = d.groupby(yr)["hub_spread"].mean()
    ok("2026's average spread sits inside 2019 to 2025's range", mn.loc[2019:2025].min() <= mn[2026] <= mn.loc[2019:2025].max())
    ok("2026's share of flat days (under $8) is at least twice any ordinary year's",
       fl[2026] >= 2 * max(fl[y] for y in (2019, 2021, 2022, 2023, 2024, 2025)) and fl[2020] > 0.25 and fl[2021] < 0.2,
       ", ".join("%d %.0f%%" % (y, 100 * fl[y]) for y in range(2019, 2027)))

    # ---------------------------------------------------------------------------------------- the asks
    cyc, lost = M.warranty_cycles(target)
    net = M.dam_net(target)
    ok("Ask A golden matches figures.json", all(round(float(cyc[m]), 2) == F["cycles"][str(m)] for m in M.OPS_MONTHS))
    ok("Ask B golden matches figures.json", all(round(float(net[m]) / 1000, 1) == F["net_k"][str(m)] for m in M.OPS_MONTHS))
    ok("Ask A every month at least 0.0008 from its rounding edge", min(edge(float(v), 0.01) for v in cyc) >= 0.0008,
       "%.4f" % min(edge(float(v), 0.01) for v in cyc))
    ok("Ask B every month at least $10 from its rounding edge", min(edge(float(v) / 1000, 0.1) for v in net) >= 0.010,
       "%.4f" % min(edge(float(v) / 1000, 0.1) for v in net))
    ok("Ask C the net per cycle at least $0.10 from its rounding edge", edge(float(net.sum()) / float(cyc.sum()), 1) >= 0.1)
    b = M.bms(target)
    m = M.meter(target)
    tw = M.test_windows(target)
    intest = pd.Series(False, index=b.index)
    for a0, z0 in tw:
        intest |= (b["local"] >= a0.replace(minute=0)) & (b["local"] < z0)
    dc = b["dc_discharge_kwh"] / 1000.0

    def by_month(vals, when):
        s = pd.Series(np.asarray(vals, float), index=when.index)
        sel = (when.dt.year == 2026) & when.dt.month.between(1, 10)
        return (s[sel].groupby(when[sel].dt.month).sum() / 400.0).reindex(range(1, 11))
    stops = {
        "the DC discharge as exported": by_month(dc.where(~intest, 0), b["local"]),
        "the same in UTC months": by_month(dc.where(~intest, 0), b["start_utc"]),
        "the TEST code alone left out": by_month(dc.where(b["operating_mode"] != "TEST", 0), b["local"]),
        "no test left out": by_month(dc, b["local"]),
        "the AC column": by_month((b["pcs_ac_discharge_kwh"] / 1000.0).where(~intest, 0), b["local"]),
    }
    dl = m.groupby("hour_utc")["mwh_delivered"].sum()
    mloc = M.local_from_utc(pd.Series(dl.index.strftime("%Y-%m-%dT%H:%M:%SZ")))
    mintest = pd.Series(False, index=mloc.index)
    for a0, z0 in tw:
        mintest |= (mloc >= a0.replace(minute=0)) & (mloc < z0)
    stops["the meter everywhere over 0.975"] = by_month((pd.Series(dl.values) / 0.975).where(~mintest, 0), mloc)
    for label, v in stops.items():
        hits = [mm for mm in range(1, 11) if round(float(v[mm]), 2) == round(float(cyc[mm]), 2)]
        if label in ("the DC discharge as exported", "the AC column", "the meter everywhere over 0.975"):
            ok("Ask A stop '%s' misses every month" % label, not hits and (v - cyc).abs().min() >= 0.012,
               "closest %.3f" % (v - cyc).abs().min())
        else:
            moved = [mm for mm in range(1, 11) if round(float(v[mm]), 2) != round(float(cyc[mm]), 2)]
            ok("Ask A stop '%s' misses every month (it carries the dropped hours too)" % label, not hits,
               "moved %s" % moved)
    drops_by_month = pd.Series([to.month for to in M.local_from_utc(pd.Series([t.strftime("%Y-%m-%dT%H:%M:%SZ") for t in lost]))]).value_counts()
    ok("Ask A primary: every month has dropped discharge hours", all(mm in drops_by_month.index for mm in range(1, 11)),
       "%d hours" % len(lost))
    ok("Ask A no dropped hour touches a capacity test", not any(intest[b["start_utc"].isin(lost)]))
    modes = b.groupby(b["local"].dt.date)["operating_mode"].agg(set)
    ok("Ask A hazard H2: the March test logged TEST, the September test MAINT",
       "TEST" in modes[W.TESTS[0]] and "MAINT" in modes[W.TESTS[1]] and "TEST" not in modes[W.TESTS[1]])
    s = pd.read_csv(M.find(target, "redrock_draw_dam_settlement_statement_lines"))
    s["od"] = pd.to_datetime(s["operating_date"])
    s = s[(s["od"].dt.year == 2026) & s["od"].dt.month.between(1, 10)]
    latest = s.sort_values("issue_date").groupby(["operating_date", "hour_ending", "repeated_hour", "charge_type"]).tail(1)

    def net_of(df, fee):
        e = df[df["charge_type"] != "DAAF"]
        f = df[df["charge_type"] == "DAAF"]
        if fee == "once":
            f = f.groupby("operating_date").head(1)
        elif fee == "none":
            f = f.iloc[0:0]
        t = pd.concat([e, f])
        return (-t.groupby(t["od"].dt.month)["amount_usd"].sum()).reindex(range(1, 11))
    bstops = {"every fee line": net_of(latest, "all"), "no fee": net_of(latest, "none"),
              "MPI's sign kept": -net}
    for label, v in bstops.items():
        hits = [mm for mm in range(1, 11) if round(float(v[mm]) / 1000, 1) == round(float(net[mm]) / 1000, 1)]
        ok("Ask B stop '%s' misses every month" % label, not hits)
    ini = s[s["statement_type"] == "INITIAL"]
    vi = net_of(ini, "once")
    ok("Ask B hazard H4: the initial statements move July only", [mm for mm in range(1, 11) if round(
        float(vi[mm]) / 1000, 1) != round(float(net[mm]) / 1000, 1)] == [7])
    fee_rows = latest[latest["charge_type"] == "DAAF"]
    ok("Ask B primary: the fee sits on every hour line of each day (23 to 25 lines)",
       fee_rows.groupby("operating_date").size().between(23, 25).all())
    ok("Ask C the net per cycle", int(round(float(net.sum()) / float(cyc.sum()))) == F["per_cycle"], str(F["per_cycle"]))
    # the separation line
    tmp = tempfile.mkdtemp(dir=TASK, suffix=".chk")
    try:
        for f in files:
            if not f.startswith(("redrock_draw_bms", "redrock_draw_revenue", "redrock_draw_dam_settlement")):
                shutil.copy(os.path.join(target, f), tmp)
        g2 = M.golden(M.daily_spreads(tmp))
        ok("separation: the main call computed without the BMS, meter and statement files is identical",
           g2["call"] == g["call"] and all(abs(g2["exp"][s_] - g["exp"][s_]) < 1e-9 for s_ in M.MENU))
        shutil.rmtree(tmp)
        os.makedirs(tmp)
        for f in files:
            if not f.startswith("mpi_dam_settlement_point_prices"):
                shutil.copy(os.path.join(target, f), tmp)
        c2, _ = M.warranty_cycles(tmp)
        n2 = M.dam_net(tmp)
        ok("separation: the asks computed without the price files are identical",
           np.allclose(c2.values, cyc.values) and np.allclose(n2.values, net.values))
        # the distractors
        for k in W.DISTRACTORS:
            shutil.rmtree(tmp)
            os.makedirs(tmp)
            for f in files:
                if f != W.FILES[k]:
                    shutil.copy(os.path.join(target, f), tmp)
            g3 = M.golden(M.daily_spreads(tmp))
            c3, _ = M.warranty_cycles(tmp)
            n3 = M.dam_net(tmp)
            ok("distractor %s: the call and every ask unchanged without it" % W.FILES[k],
               g3["call"] == g["call"] and np.allclose(c3.values, cyc.values) and np.allclose(n3.values, net.values))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    # ------------------------------------------------------------------------------- input gates and hygiene
    fmts = {f.rsplit(".", 1)[1] for f in files}
    rows = {f: sum(1 for _ in open(os.path.join(target, f), encoding="utf-8")) - 1 for f in files if f.endswith(".csv")}
    ok("input gates: 10 or more files, 3 or more formats, a file of 25,000 rows or more",
       len(files) >= 10 and len(fmts) >= 3 and max(rows.values()) >= 25000, "%d files, %s, largest %d rows" % (
           len(files), sorted(fmts), max(rows.values())))
    ok("two distractors named in metadata.json, none in a file name",
       len(meta["distractor_files"]) == 2 and all(f in files for f in meta["distractor_files"]) and
       not any("distract" in f for f in files))
    man = pd.read_csv(os.path.join(target, W.FILES["manifest"]))
    ok("the manifest lists every shipped file but itself", set(man["file"]) == set(files) - {W.FILES["manifest"]})
    import openpyxl
    dic = openpyxl.load_workbook(os.path.join(target, W.FILES["dict"]))["Field definitions"]
    named = {c.value for c in dic["A"] if c.value}
    ok("the dictionary names every file in its scope", (set(files) - {W.FILES["dict"], W.FILES["manifest"]}) <= named)
    texts = [open(os.path.join(target, f), encoding="utf-8").read() for f in files if f.endswith(".csv")]
    texts += list(doc_text.values())
    texts += [" ".join(str(c.value) for row in dic.iter_rows() for c in row if c.value)]
    texts += [open(os.path.join(TASK, n), encoding="utf-8").read() for n in ("prompt.txt", "submission.md", "DESIGN_NOTE.md")]
    ok("no em or en dash in the pack, the prompt, the submission or the design note",
       not any("—" in t or "–" in t for t in texts))
    people = re.findall(r"[A-Z][a-z]+ [A-Z][a-z]+", " ".join(doc_text.values()))
    ok("the only people named are the drawn personas", W.PEOPLE["asset"] in pol and not any(
        p in ("John Doe", "Jane Doe") for p in people))
    ok("pull dates hold: no record after its file's pull",
       d["day"].max() <= W.ARCHIVE_PULL and b["start_utc"].max() < dt.datetime(2026, 11, 3) and
       pd.to_datetime(s["issue_date"]).max() <= pd.Timestamp(W.SETTLE_PULL))
    ok("realism: no column constant in the BMS, meter or statement extracts", all(
        pd.read_csv(os.path.join(target, W.FILES[k]))[c].nunique() > 1 for k in ("meter",) for c in
        ("mwh_delivered", "mwh_received")) and b["operating_mode"].nunique() >= 5 and s["charge_type"].nunique() == 3)
    ok("realism: the day-ahead awards vary from day to day", latest[latest["charge_type"] == "DAES"][
        "bill_determinant_mwh"].nunique() >= 6)
    ok("realism: no monthly cycle count a whole number and no round total",
       all(abs(float(v) - round(float(v))) > 0.004 for v in cyc) and abs(float(net.sum()) % 1000) > 1)
    ok("realism: the node's spread follows the hub's with day-to-day variation (correlation under 0.995)",
       np.corrcoef(ov["hub_spread"], ov["node_spread"])[0, 1] < 0.995)
    for f in files:
        if f.endswith(".pdf"):
            raw = open(os.path.join(target, f), "rb").read()
            ok("container: %s carries no writer name" % f, b"ReportLab" not in raw and b"reportlab" not in raw)
        if f.endswith(".xlsx"):
            with zipfile.ZipFile(os.path.join(target, f)) as z:
                core = z.read("docProps/core.xml").decode()
            ok("container: %s stamped with the fiction's people and dates" % f, "openpyxl" not in core and
               "2026-11" in core)

    # ----------------------------------------------------------------------------------------- the golden files
    out = pd.read_csv(os.path.join(golden, W.GOLDEN["daily"]))
    ok("the conformed file: one row per point and operating day, the node under one name",
       len(out) == 2863 + 336 and set(out["settlement_point"]) == {W.HUB, W.NODE} and
       not out.duplicated(["operating_date", "settlement_point"]).any())
    wb = openpyxl.load_workbook(os.path.join(golden, W.GOLDEN["workbook"]))
    ok("the workbook's sheets", wb.sheetnames == ["2027 threshold", "2026 operation", "Cycle chart", "Sources"])
    ok("the workbook's chart", len(wb["Cycle chart"]._charts) == 1)
    ok("figures.json carries the call and the eight counts", F["call"] == g["call"] and
       all(F["exp"][str(s_)] == round(g["exp"][s_]) for s_ in M.MENU))
    return report()


def report():
    bad = [r for r in RESULTS if not r[1]]
    for label, good, detail in RESULTS:
        print("%s  %s%s" % ("PASS" if good else "FAIL", label, ("  [" + detail + "]") if detail else ""))
    print("checks: %d run, %d failed" % (len(RESULTS), len(bad)))
    return not bad


if __name__ == "__main__":
    good = run(os.path.join(TASK, "target"), os.path.join(TASK, "golden"))
    raise SystemExit(0 if good else 1)
