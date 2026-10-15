"""Independent verifier for task461: reads only the shipped bundle (target/) on its own code path (the csv module,
its own clock handling and its own dropped-hour test) and recomputes the call, the expected 2027 cycling days at each
setting, the overlap's ratio, the default path's and the premium rung's settings, the 2026 Full Cycles by month, the
day-ahead net by month and per cycle; then compares them with generator/figures.json and the golden files.

    python3 task461/generator/verify.py [target] [golden]
"""
import calendar
import csv
import datetime as dt
import json
import os
import re
import sys
from collections import defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
MENU = [5, 10, 15, 20, 25, 30, 35, 40]


def find(target, stem):
    hits = [f for f in os.listdir(target) if f.startswith(stem)]
    assert len(hits) == 1, (stem, hits)
    return os.path.join(target, hits[0])


def parse_day(s):
    if "/" in s:
        m, d, y = s.split("/")
        return dt.date(int(y), int(m), int(d))
    return dt.date.fromisoformat(s)


def spreads(path, rename):
    by = defaultdict(list)
    with open(path, newline="") as f:
        for r in csv.DictReader(f):
            by[(rename.get(r["SettlementPoint"], r["SettlementPoint"]), parse_day(r["DeliveryDate"]))].append(
                float(r["SettlementPointPrice"]))
    out = {}
    for k, v in by.items():
        s = sorted(v)
        out[k] = sum(s[-4:]) / 4.0 - sum(s[:4]) / 4.0
    return out


def expected(vals):
    n = len(vals)
    return {s: 365.0 * sum(1 for v in vals if v >= s) / n for s in MENU}


def pick(exp):
    return min(s for s in MENU if exp[s] <= 300)


def dst(y):
    m = dt.date(y, 3, 1)
    spring = m + dt.timedelta(days=(6 - m.weekday()) % 7 + 7)
    n = dt.date(y, 11, 1)
    fall = n + dt.timedelta(days=(6 - n.weekday()) % 7)
    return dt.datetime(spring.year, spring.month, spring.day, 8), dt.datetime(fall.year, fall.month, fall.day, 7)


def to_local(t):
    a, b = dst(t.year)
    return t - dt.timedelta(hours=5 if a <= t < b else 6)


def main(target=None, golden=None):
    target = target or os.path.join(TASK, "target")
    golden = golden or os.path.join(TASK, "golden")
    F = json.load(open(os.path.join(HERE, "figures.json")))
    bad = []

    def same(label, a, b):
        if a != b:
            bad.append("%s: verifier %s, generator %s" % (label, a, b))

    # the settlement point list: the node's names and the last day priced at the hub
    with open(find(target, "mpi_settlement_point_list")) as f:
        pts = list(csv.DictReader(f))
    rename = {p["predecessor"]: p["settlement_point"] for p in pts if p["predecessor"]}
    node_name = [p["settlement_point"] for p in pts if p["predecessor"] and p["point_type"] == "RN" and
                 "ESS" in p["settlement_point"] and p["predecessor"] in rename and rename[p["predecessor"]] == p["settlement_point"]
                 and p["predecessor"].startswith("RDRW")][0]
    note = [p["note"] for p in pts if p["settlement_point"].startswith("RDRW")][0]
    last_at_hub = dt.date.fromisoformat(re.search(r"through operating day (\d{4}-\d{2}-\d{2})", note).group(1))
    hub = spreads(find(target, "mpi_dam_settlement_point_prices_north_hub"), {})
    node = spreads(find(target, "mpi_dam_settlement_point_prices_redrock"), rename)
    hd = {d: v for (p, d), v in hub.items()}
    nd = {d: v for (p, d), v in node.items() if p == node_name}
    days = sorted(hd)
    ov = [d for d in nd if d > last_at_hub]
    k = sum(nd[d] for d in ov) / sum(hd[d] for d in ov)
    ser = [nd[d] if d in nd and d > last_at_hub else k * hd[d] for d in days]
    exp = expected(ser)
    same("call", pick(exp), F["call"])
    for s in MENU:
        same("expected days at $%d" % s, round(exp[s]), F["exp"][str(s)])
    same("ratio", round(k, 2), F["k"])
    same("overlap days", len(ov), F["overlap_days"])
    r0 = expected([nd[d] for d in nd])
    same("R0 setting", pick(r0), F["rungs"]["R0 the node's own record"]["call"])
    prem = sum(nd[d] - hd[d] for d in ov) / len(ov)
    same("R1 setting", pick(expected([hd[d] + prem for d in days])),
         F["rungs"]["R1 the hub plus the node's average spread premium"]["call"])
    same("hub alone setting", pick(expected([hd[d] for d in days])), F["rungs"]["Rx the hub's record alone"]["call"])

    # Full Cycles: the BMS's DC discharge; where the BMS shows none and the meter shows the battery delivering, the
    # meter over 0.975 (warranty 4.3); the capacity tests' windows from Cordell's reports left out
    from pypdf import PdfReader
    txt = re.sub(r"\s+", " ", " ".join(p.extract_text() for p in PdfReader(find(target, "cordell_capacity_test")).pages))
    wins = []
    for m in re.finditer(r"Discharge: (\d{4}-\d{2}-\d{2}) (\d{2}):(\d{2}) to (\d{2}):(\d{2})", txt):
        d0 = dt.date.fromisoformat(m.group(1))
        wins.append((dt.datetime(d0.year, d0.month, d0.day, int(m.group(2))),
                     dt.datetime(d0.year, d0.month, d0.day, int(m.group(4)), int(m.group(5)))))
    meter = defaultdict(float)
    with open(find(target, "redrock_draw_revenue_meter")) as f:
        for r in csv.DictReader(f):
            end = dt.datetime.strptime(r["interval_ending_utc"], "%Y-%m-%dT%H:%M:%SZ")
            meter[(end - dt.timedelta(minutes=15)).replace(minute=0)] += float(r["mwh_delivered"])
    cyc = defaultdict(float)
    prev = None
    with open(find(target, "redrock_draw_bms_hourly_export")) as f:
        for r in csv.DictReader(f):
            t = dt.datetime.strptime(r["interval_start_utc"], "%Y-%m-%dT%H:%M:%SZ")
            loc = to_local(t)
            dc = float(r["dc_discharge_kwh"]) / 1000.0
            gap = (float(r["dc_discharge_kwh"]) == 0 and float(r["dc_charge_kwh"]) == 0 and prev is not None and
                   r["soc_pct"] == prev["soc_pct"] and r["operating_mode"] == prev["operating_mode"])
            e = meter.get(t, 0.0) / 0.975 if gap and meter.get(t, 0.0) > 0 else dc
            if any(a <= loc < b for a, b in wins):
                e = 0.0
            if loc.year == 2026 and loc.month <= 10:
                cyc[loc.month] += e / 400.0
            prev = r
    for m in range(1, 11):
        same("Full Cycles %s" % calendar.month_abbr[m], round(cyc[m], 2), F["cycles"][str(m)])
    same("Full Cycles total", round(sum(cyc.values()), 2), F["cycles_total"])

    # the day-ahead net: the latest line per hour and charge type, MPI's sign reversed, the fee once a day
    latest = {}
    with open(find(target, "redrock_draw_dam_settlement_statement_lines")) as f:
        for r in csv.DictReader(f):
            od = dt.date.fromisoformat(r["operating_date"])
            if od.year != 2026 or od.month > 10:
                continue
            key = (od, r["hour_ending"], r["repeated_hour"], r["charge_type"])
            if key not in latest or r["issue_date"] > latest[key]["issue_date"]:
                latest[key] = r
    net = defaultdict(float)
    fee_days = set()
    for (od, he, rep, ct), r in latest.items():
        if ct == "DAAF":
            if od in fee_days:
                continue
            fee_days.add(od)
        net[od.month] -= float(r["amount_usd"])
    for m in range(1, 11):
        same("net %s" % calendar.month_abbr[m], round(net[m] / 1000.0, 1), F["net_k"][str(m)])
    same("net total", round(sum(net.values()) / 1000.0, 1), F["net_total_k"])
    same("net per cycle", int(round(sum(net.values()) / sum(cyc.values()))), F["per_cycle"])

    # the golden files
    with open(os.path.join(golden, "redrock_draw_daily_spreads_conformed.csv")) as f:
        rows = list(csv.DictReader(f))
    same("conformed rows", len(rows), F["csv_rows"])
    same("conformed points", sorted({r["settlement_point"] for r in rows}), sorted(["MP_NORTH_HUB", node_name]))
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(golden, "redrock_draw_2027_cycling_threshold.xlsx"))
    ws = wb["2027 threshold"]
    table = {int(ws.cell(row=r, column=1).value): ws.cell(row=r, column=2).value for r in range(5, 13)}
    for s in MENU:
        same("workbook days at $%d" % s, table[s], round(exp[s]))
    ws2 = wb["2026 operation"]
    for i, m in enumerate(range(1, 11)):
        same("workbook cycles %d" % m, ws2.cell(row=5 + i, column=2).value, round(cyc[m], 2))
        same("workbook net %d" % m, ws2.cell(row=5 + i, column=5).value, round(net[m] / 1000.0, 1))
    if bad:
        print("VERIFY FAILED:\n  " + "\n  ".join(bad))
        raise SystemExit(1)
    print("verify: independent recomputation agrees: $%d (%d expected days; $%d %d), ratio %.2f over %d days; R0 $%d, "
          "R1 $%d, the hub alone $%d; 2026 Full Cycles %.2f, day-ahead net $%.1fk, $%s per cycle; the conformed file's "
          "%d rows and the workbook's tables" % (
              F["call"], round(exp[F["call"]]), F["below"], round(exp[F["below"]]), k, len(ov), pick(r0),
              F["rungs"]["R1 the hub plus the node's average spread premium"]["call"],
              F["rungs"]["Rx the hub's record alone"]["call"], sum(cyc.values()), sum(net.values()) / 1000.0,
              format(F["per_cycle"], ","), len(rows)))


if __name__ == "__main__":
    main(*(sys.argv[1:3]))
