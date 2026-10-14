"""The independent verifier: reads only the shipped bundle (target/) and the goldens, on its own code path (its own
parsing of the PDFs' tables, the roster and both providers' exports; each operator's probability by Genz's
algorithm rather than the generator's quasi-Monte Carlo), and recomputes the call, every configuration's
accommodation, each option's maintenance, bus-years and cost per bus-year, and each configuration's twelve-year cost."""
import csv
import datetime as dt
import math
import os
import re
import sys
from collections import defaultdict

import numpy as np
import openpyxl
from scipy import stats


def pdf_text(path):
    from pypdf import PdfReader
    return re.sub(r"\s+", " ", " ".join(p.extract_text() for p in PdfReader(path).pages))


def find(folder, prefix):
    hits = [f for f in os.listdir(folder) if f.startswith(prefix)]
    assert len(hits) == 1, (prefix, hits)
    return os.path.join(folder, hits[0])


NAMES = ["Stature (mm)", "Weight (kg)", "Sitting height (mm)", "Eye height, sitting (mm)",
         "Shoulder height, sitting (mm)", "Popliteal height (mm)", "Knee height, sitting (mm)",
         "Buttock-knee length (mm)", "Buttock-popliteal length (mm)", "Thumb-tip reach (mm)",
         "Hip breadth, sitting (mm)", "Abdominal extension depth, sitting (mm)"]
FIT = [5, 7, 3, 9, 11]          # popliteal, buttock-knee, eye height sitting, thumb-tip reach, abdominal depth
GIVEN = [0, 1]                  # stature, weight


def survey(path):
    t = pdf_text(path)
    t1, rest = t.split("Table 1.", 1)[1].split("Table 2.", 1)
    t2, rest = rest.split("Table 3.", 1)
    t3, t4 = rest.split("Table 4.", 1)
    out = {}
    for sx, tab, corr in (("M", t1, t3), ("F", t2, t4)):
        mu, sd = [], []
        for nm in NAMES:
            m = re.search(re.escape(nm) + r" ([\d.]+) ([\d.]+)", tab)
            mu.append(float(m.group(1)))
            sd.append(float(m.group(2)))
        codes = re.search(r"working (?:men|women) (S W .*?AED)", corr).group(1).split()
        assert len(codes) == 12, codes
        body = corr.split(" ".join(codes), 1)[1]
        r = np.eye(12)
        for i, c in enumerate(codes):
            m = re.search(r"\b%s ((?:-?\d\.\d\d ){%d}1\.00)" % (re.escape(c), i), body)
            vals = [float(v) for v in m.group(1).split()]
            for j in range(i):
                r[i, j] = r[j, i] = vals[j]
        out[sx] = (np.array(mu), np.array(sd), r)
    return out


def envelopes(path):
    t = pdf_text(path)
    prices = {c: int(p.replace(",", "")) for c, p in re.findall(r"\b([A-H]) [A-Z][^$]*?\$([\d,]+)", t.split("4 Fit envelopes")[0])}
    env = {}
    cell = r"(no limit|[\d,]+) to (no limit|[\d,]+)"
    for m in re.finditer(r"\b([A-H]) " + " ".join([cell] * 5), t):
        vals = [None if v == "no limit" else int(v.replace(",", "")) for v in m.groups()[1:]]
        env[m.group(1)] = [(vals[2 * k], vals[2 * k + 1]) for k in range(5)]
    assert len(env) == 8 and len(prices) == 8, (env, prices)
    return env, prices


def main(target, golden):
    sv = survey(find(target, "brennard"))
    env, prices = envelopes(find(target, "halvard_fit_report"))
    std = pdf_text(find(target, "tvrta_standard"))
    line = 0.95 if "accommodation is at least 95 per cent" in std else None
    assert line and "whatever their duty status" in std

    # the counted operators and their latest DOT physical in either export
    wb = openpyxl.load_workbook(find(target, "hr_roster"), read_only=True)
    ops = {}
    for r in wb["Roster"].iter_rows(min_row=6, values_only=True):
        if r[1] == "BUS OPERATOR":
            ops[int(r[0])] = r[5]
    latest = {}

    def keep(emp, d, h_in, w_lb):
        if emp in ops and (emp not in latest or d > latest[emp][0]):
            latest[emp] = (d, h_in, w_lb)
    for r in csv.DictReader(open(find(target, "ostrander"))):
        if r["SERVICE"] == "DOT PHYSICAL":
            ft, inch = re.match(r"(\d+)' (\d+)\"", r["HEIGHT"]).groups()
            keep(int(r["EMP#"]), dt.datetime.strptime(r["DOS"], "%m/%d/%Y").date(), int(ft) * 12 + int(inch),
                 int(r["WEIGHT LB"]))
    for r in csv.DictReader(open(find(target, "fernhill"))):
        if r["visit_type"] == "DOTPE":
            keep(int(r["employer_ref"].split("-")[1]), dt.date.fromisoformat(r["visit_date"]), float(r["ht_in"]),
                 int(r["wt_lb"]))
    assert set(latest) == set(ops), "an operator without a physical"

    # how far the operators Halvard measured sit above the survey's eye height for their own stature and weight
    shift = {}
    for sx in ("M", "F"):
        mu, sd, r = sv[sx]
        cov = np.outer(sd, sd) * r
        k = cov[np.ix_(FIT, GIVEN)] @ np.linalg.inv(cov[np.ix_(GIVEN, GIVEN)])
        dev = []
        for t in csv.DictReader(open(find(target, "halvard_operator_measurements"))):
            if t["sex"] != sx:
                continue
            m = mu[FIT] + k @ (np.array([float(t["stature_mm"]), float(t["weight_kg"])]) - mu[GIVEN])
            dev.append(float(t["eye_height_sitting_mm"]) - m[2])
        shift[sx] = np.array([0.0, 0.0, float(np.mean(dev)), 0.0, 0.0])

    acc = {}
    for c in "ABCDEFGH":
        lo = np.array([-np.inf if env[c][k][0] is None else env[c][k][0] for k in range(5)], float)
        hi = np.array([np.inf if env[c][k][1] is None else env[c][k][1] for k in range(5)], float)
        tot = 0.0
        for sx in ("M", "F"):
            mu, sd, r = sv[sx]
            cov = np.outer(sd, sd) * r
            a, b = FIT, GIVEN
            k = cov[np.ix_(a, b)] @ np.linalg.inv(cov[np.ix_(b, b)])
            resid = cov[np.ix_(a, a)] - k @ cov[np.ix_(b, a)]
            mv = stats.multivariate_normal(np.zeros(5), resid, abseps=1e-6, releps=1e-6, maxpts=200000)
            for emp, s in ops.items():
                if s != sx:
                    continue
                _, h, w = latest[emp]
                m = mu[a] + k @ (np.array([h * 25.4, w * 0.45359237]) - mu[b]) + shift[sx]
                tot += mv.cdf(hi - m, lower_limit=lo - m, rng=np.random.default_rng(emp))
        acc[c] = tot / len(ops)
    order = sorted(prices, key=prices.get)
    call = next(c for c in order if acc[c] >= line)

    # the ask: option maintenance from the work orders, bus-years from the fleet, disposals and retrofits
    code = {"PDL-ADJ": "PDL", "SEAT-RSR": "RSR", "STR-TEL": "TEL", "SEAT-TRK": "TRK", "HDR-VIS": "HDR"}
    cost = defaultdict(float)
    seen = set()
    for r in csv.DictReader(open(find(target, "operator_area_work_orders"))):
        if r["component"] not in code or r["bill_to"] != "TVRTA" or not ("2021" <= r["closed"][:4] <= "2025"):
            continue
        cost[code[r["component"]]] += float(r["part_cost"])
        if r["wo_number"] not in seen:
            seen.add(r["wo_number"])
            cost[code[r["component"]]] += float(r["labor_cost"])
    buses = []
    for r in csv.DictReader(open(find(target, "fleet_roster"))):
        buses.append((dt.date.fromisoformat(r["accepted"]), None, set(r["operator_area_options"].split()),
                      r["builder_serial"]))
    for r in csv.DictReader(open(find(target, "bus_disposals"))):
        buses.append((dt.date.fromisoformat(r["accepted"]), dt.date.fromisoformat(r["disposed"]),
                      set(r["operator_area_options"].split()), r["builder_serial"]))
    fitted = {}
    for r in csv.DictReader(open(find(target, "retrofit_campaigns"))):
        if r["operator_area_option"]:
            fitted[(r["builder_serial"], r["operator_area_option"])] = dt.date.fromisoformat(r["completed"])
    months = defaultdict(int)
    for y in range(2021, 2026):
        for mo in range(1, 13):
            d = dt.date(y, mo, 1)
            for acc_d, gone, opts, ser in buses:
                if acc_d <= d and (gone is None or d < gone):
                    for o in opts:
                        if fitted.get((ser, o), acc_d) <= d:
                            months[o] += 1
    rate = {o: cost[o] / (months[o] / 12) for o in code.values()}
    opt_of = {"Power-adjustable pedal set": "PDL", "Extended-range seat riser": "RSR",
              "Tilt-telescope steering column": "TEL", "Deep-travel seat track": "TRK", "Raised header and visor": "HDR"}
    t = pdf_text(find(target, "halvard_fit_report")).split("4 Fit envelopes")[0]
    twelve = {}
    for c in "ABCDEFGH":
        row = re.search(r"\b%s (.*?)\$([\d,]+)" % c, t).group(1)
        twelve[c] = prices[c] + 12 * sum(round(rate[v]) for k, v in opt_of.items() if k in row)

    # the goldens
    gw = openpyxl.load_workbook(os.path.join(golden, "workstation_configurations_2027.xlsx"))
    ws = gw["Configurations"]
    assert "Specify Configuration %s" % call in ws["A2"].value, ws["A2"].value
    for r in ws.iter_rows(min_row=6, max_row=13, values_only=True):
        c = r[0]
        assert round(100 * acc[c], 1) == round(100 * r[3], 1), ("accommodation", c, acc[c], r[3])
        assert abs(acc[c] - r[3]) < 0.0006, ("accommodation", c, acc[c], r[3])
        assert r[2] == prices[c] and r[5] == round(twelve[c]), ("price or twelve-year", c, r, twelve[c])
    mt = gw["Option maintenance"]
    for r in mt.iter_rows(min_row=5, max_row=9, values_only=True):
        o = opt_of[r[0]]
        assert r[2] == round(cost[o]) and r[3] == round(months[o] / 12, 1) and r[4] == round(rate[o]), (r, o)
    print("verify: from target/ alone, %d operators, call %s (%s); accommodation %s; option rates %s; twelve-year %s; "
          "goldens agree" % (len(ops), call, "%.2f%%" % (100 * acc[call]),
                             ", ".join("%s %.2f" % (c, 100 * acc[c]) for c in order),
                             ", ".join("%s $%d/%.1f" % (o, round(rate[o]), months[o] / 12) for o in code.values()),
                             ", ".join("%s $%s" % (c, "{:,}".format(round(twelve[c]))) for c in order)))


if __name__ == "__main__":
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "target"),
         sys.argv[2] if len(sys.argv) > 2 else os.path.join(here, "golden"))
