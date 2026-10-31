"""The pack's data files, each as its source system exports it: Human Resources' roster, the two occupational health
providers' encounter exports, Fleet Maintenance's fleet roster, disposals, retrofit campaigns and operator-area work
orders, and two files kept beside them (the 2023 operator comfort survey and the workers' compensation claims)."""
import csv
import datetime as dt

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

import world as W

ROSTER = "hr_roster_transportation_maintenance_2026-09-30.xlsx"
OSTRANDER = "ostrander_encounters_tvrta_2014-01_2021-06.csv"
FERNHILL = "fernhill_visits_tvrta_2021-01_2026-09.csv"
FLEET = "fleet_roster_2026-09-30.csv"
DISPOSALS = "bus_disposals_2019-2026.csv"
RETROFITS = "retrofit_campaigns_2019-2026.csv"
WORK_ORDERS = "operator_area_work_orders_2019-01_2026-09.csv"
COMFORT = "operator_comfort_survey_2023.xlsx"
CLAIMS = "workers_comp_claims_transportation_2019-2026.csv"
TRIAL = "halvard_operator_measurements_2026-08.csv"

OPT_CODE = {"P": "PDL", "R": "RSR", "T": "TEL", "D": "TRK", "V": "HDR"}
OST_SERVICE = {"PHYS": "DOT PHYSICAL", "RECERT": "DOT PHYSICAL - RECERT", "DRUG-RAND": "DRUG SCREEN - RANDOM",
               "DRUG-PRE": "DRUG SCREEN - PRE-EMPLOYMENT", "DRUG-PA": "DRUG SCREEN - POST ACCIDENT",
               "ALC-RAND": "BREATH ALCOHOL - RANDOM"}
FH_TYPE = {"PHYS": "DOTPE", "RECERT": "DOTRC", "DRUG-RAND": "UDS-RAND", "DRUG-PRE": "UDS-PREEMP",
           "DRUG-PA": "UDS-POSTACC", "ALC-RAND": "BAT-RAND"}
FH_DESC = {"PHYS": "DOT physical examination", "RECERT": "DOT recertification, short certificate",
           "DRUG-RAND": "Urine drug screen, random", "DRUG-PRE": "Urine drug screen, pre-employment",
           "DRUG-PA": "Urine drug screen, post-accident", "ALC-RAND": "Breath alcohol test, random"}


def us(d):
    return d.strftime("%m/%d/%Y")


def emp_text(n):
    return "%06d" % n


def feet_inches(h_in):
    k = int(round(h_in))
    return "%d' %02d\"" % (k // 12, k % 12)


def cert_expiry(e):
    if e["type"] == "PHYS":
        months = e["months"]
        base = e["date"]
    elif e["type"] == "RECERT":
        months = 12
        base = e["of"]
    else:
        return None
    y, m = base.year + (base.month - 1 + months) // 12, (base.month - 1 + months) % 12 + 1
    last = [31, 29 if y % 4 == 0 else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][m - 1]
    return dt.date(y, m, min(base.day, last))


def ostrander_numbers(enc):
    """Ostrander's encounter numbers: year and month of service, then the clinic's running count in that month."""
    seq, out = {}, {}
    for i, e in enumerate(enc):
        if e["date"] >= W.CHANGEOVER:
            continue
        key = e["date"].strftime("%y%m")
        seq[key] = seq.get(key, 0) + 1 + W.hint("osgap", i) % 4
        out[i] = "%s-%05d" % (key, seq[key])
    return out


def write_roster(path, people):
    cur = [p for p in people if p["left"] is None]
    cur.sort(key=lambda p: (W.DIVISIONS.index(p["division"]), p["job"] != "BUS OPERATOR", p["job"], p["emp"]))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Roster"
    ws["A1"] = W.AUTH
    ws["A1"].font = Font(bold=True)
    ws["A2"] = "Human Resources: employees in Transportation and Maintenance job classes"
    ws["A3"] = "As of 09/30/2026"
    ws.append([])
    heads = ["Emp No", "Job Class", "FT/PT", "Division", "Hire Date", "Sex", "Duty Status"]
    ws.append(heads)
    for c in ws[5]:
        c.font = Font(bold=True)
    for p in cur:
        ft = ("PT" if p.get("pt") else "FT")
        ws.append([emp_text(p["emp"]), p["job"], ft, p["division"].upper(), p["hired"], p["sex"], p["status"]])
        ws.cell(row=ws.max_row, column=5).number_format = "mm/dd/yyyy"
    for i, wd in enumerate((10, 30, 7, 13, 12, 6, 15), 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = "A5:G%d" % ws.max_row
    wb.save(path)
    return len(cur)


def write_ostrander(path, enc, nums):
    n = 0
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["ACCT", "ENC #", "EMP#", "DOS", "SERVICE", "HEIGHT (NO SHOES)", "WEIGHT LB", "BP", "CERT", "CERT EXPIRES"])
        for i, e in enumerate(enc):
            if e["date"] >= W.CHANGEOVER:
                continue
            phys = e["type"] == "PHYS"
            bp = "%d/%d" % e["bp"] if phys else ("%d/%d" % (128 + W.hint("rc", i) % 14, 80 + W.hint("rd", i) % 9)
                                                  if e["type"] == "RECERT" else "")
            cert = ("%d MONTHS" % e["months"] if phys else "9 MONTHS" if e["type"] == "RECERT" else "")
            exp = cert_expiry(e)
            w.writerow(["TVRTA-01", nums[i], str(e["emp"]), us(e["date"]), OST_SERVICE[e["type"]],
                        feet_inches(e["h_in"]) if phys else "", e["w_lb"] if phys else "", bp, cert,
                        us(exp) if exp else ""])
            n += 1
    return n


def write_fernhill(path, enc, nums):
    n = 0
    seq = {}
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["visit_id", "employer_ref", "visit_date", "visit_type", "visit_desc", "ht_in_no_shoes", "wt_lb",
                    "bp_sys", "bp_dia",
                    "cert_months", "cert_expires", "legacy_ref"])
        for i, e in enumerate(enc):
            if e["date"] < W.TAKEN_OVER:
                continue
            taken = e["date"] < W.CHANGEOVER
            key = "%02d%02d" % (e["date"].year % 100, e["date"].month)
            seq[key] = seq.get(key, 0) + 1
            vid = "FH-%s-%04d" % (key, seq[key] + (500 if taken else 0))
            phys = e["type"] == "PHYS"
            if phys:
                sys_, dia = e["bp"]
            elif e["type"] == "RECERT":
                sys_, dia = 128 + W.hint("rc", i) % 14, 80 + W.hint("rd", i) % 9
            else:
                sys_ = dia = ""
            months = e["months"] if phys else (9 if e["type"] == "RECERT" else "")
            exp = cert_expiry(e)
            w.writerow([vid, "TVT-%d" % e["emp"], e["date"].isoformat(), FH_TYPE[e["type"]], FH_DESC[e["type"]],
                        "%.1f" % e["h_in"] if phys else "", e["w_lb"] if phys else "", sys_, dia, months,
                        exp.isoformat() if exp else "", nums.get(i, "") if taken else ""])
            n += 1
    return n


def opt_text(opts):
    return " ".join(OPT_CODE[o] for o in W.OPTION_ORDER if o in opts)


def write_fleet(path, buses):
    cur = sorted((b for b in buses if b["disposed"] is None), key=lambda b: b["no"])
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["fleet_no", "model_year", "builder_serial", "model", "length_ft", "division", "accepted",
                    "operator_area_options", "status"])
        for b in cur:
            st = "ACTIVE" if W.hint("spare", b["no"], b["series"]) % 19 else "SPARE"
            w.writerow([b["no"], b["series"], b["serial"], "XE40" if b["series"] >= 2021 else "LF40",
                        40, b["division"].upper(), b["accepted"].isoformat(), opt_text(W.options_on(b, W.AS_OF)), st])
    return len(cur)


def write_disposals(path, buses):
    gone = sorted((b for b in buses if b["disposed"] is not None), key=lambda b: (b["disposed"], b["no"]))
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["fleet_no", "model_year", "builder_serial", "accepted", "disposed", "operator_area_options",
                    "disposition"])
        for b in gone:
            w.writerow([b["no"], b["series"], b["serial"], b["accepted"].isoformat(), b["disposed"].isoformat(),
                        opt_text(W.options_on(b, b["disposed"])),
                        "AUCTION" if W.hint("disp", b["no"], b["series"]) % 5 else "SCRAP"])
    return len(gone)


def write_retrofits(path, buses):
    rows = []
    for b in buses:
        for o, when in b["retrofit"].items():
            rows.append((when, b, "24-03", "Power-adjustable pedal set, retrofit", OPT_CODE[o]))
        if b["series"] in (2014, 2015, 2016) and W.hint("barr", b["no"]) % 3 == 0:
            when = dt.date(2022, 9, 12) + dt.timedelta(days=W.hint("barrd", b["no"]) % 60)
            rows.append((when, b, "22-05", "Operator barrier door, full-height replacement", ""))
    rows.sort(key=lambda r: (r[0], r[1]["no"]))
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["campaign", "description", "fleet_no", "builder_serial", "operator_area_option", "completed"])
        for when, b, camp, desc, code in rows:
            w.writerow([camp, desc, b["no"], b["serial"], code, when.isoformat()])
    return len(rows)


def write_work_orders(path, wos):
    n = 0
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["wo_number", "line", "fleet_no", "component", "component_desc", "opened", "closed", "reason",
                    "bill_to", "part_number", "part_desc", "qty", "unit_cost", "part_cost", "labor_hours",
                    "labor_cost"])
        for wo in sorted(wos, key=lambda o: (o["opened"], o["wo"])):
            labour = round(wo["hours"] * wo["rate"], 2)
            for k, ln in enumerate(wo["lines"], 1):
                w.writerow([wo["wo"], k, wo["no"], wo["comp"], W.COMPONENTS[wo["comp"]][0], wo["opened"].isoformat(),
                            wo["closed"].isoformat(), wo["reason"], wo["bill"], ln["part"], ln["desc"], ln["qty"],
                            "%.2f" % ln["unit"], "%.2f" % round(ln["qty"] * ln["unit"], 2), "%.1f" % wo["hours"],
                            "%.2f" % labour])
                n += 1
    return n


def write_comfort(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Responses"
    heads = ["Response ID", "Submitted", "Division", "Years as operator", "Bus series driven most",
             "Seat comfort (1-5)", "Reach to controls (1-5)", "Steering wheel clearance (1-5)",
             "Visibility over the dash (1-5)", "Mirror adjustment (1-5)"]
    ws.append(heads)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append([r["id"], r["date"], r["division"].upper(), r["years"], r["series"]] + r["ratings"])
        ws.cell(row=ws.max_row, column=2).number_format = "mm/dd/yyyy"
    for i, wd in enumerate((11, 11, 12, 10, 12, 10, 10, 12, 12, 11), 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = "A2"
    wb.save(path)
    return len(rows)


def write_claims(path, rows):
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["claim_no", "emp_no", "date_of_injury", "job_class", "body_part", "cause", "lost_days",
                    "paid_to_date", "status"])
        for r in rows:
            w.writerow([r["claim"], emp_text(r["emp"]), r["date"].isoformat(), r["job"], r["body"], r["cause"],
                        r["lost"], "%.2f" % r["paid"], r["status"]])
    return len(rows)


def write_trial(path, trial):
    """Halvard's measurement sheet: one row per operator measured, eye height in the survey's posture."""
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["subject", "session_date", "sex", "stature_mm", "weight_kg", "eye_height_sitting_mm"])
        for t in trial:
            w.writerow([t["participant"], t["day"].isoformat(), t["sex"], "%.0f" % t["s"], "%.1f" % t["w"],
                        "%.0f" % t["seh"]])
    return len(trial)
