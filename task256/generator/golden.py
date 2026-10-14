"""The golden deliverables: Fleet Engineering's configuration workbook for the Board and the chart of each
configuration's share of the operators it fits."""
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt                       # noqa: E402
import openpyxl                                       # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter          # noqa: E402

import world as W                                     # noqa: E402

WORKBOOK = "workstation_configurations_2027.xlsx"
CHART = "operator_accommodation_2027.svg"
NAVY, TEAL, PALE, GREYTXT = "1F3A5F", "2E7D6B", "E8F1EE", "5D646C"


def write_workbook(path, fig):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Configurations"
    call = fig["call"]
    ops = fig["operators"]
    ws["A1"] = W.AUTH
    ws["A1"].font = Font(bold=True, color=NAVY, size=12)
    ws["A2"] = ("Specify Configuration %s for the 2027 order: the lowest-priced configuration that fits at least 95 per "
                "cent of our bus operators (%.1f per cent, ${:,} per bus)".format(call["price"])
                % (call["config"], call["acc_pct"]))
    ws["A2"].font = Font(bold=True, size=11)
    ws["A3"] = ("Standard %s, section 3; the {:,} bus operators on the roster of 30 September 2026 at their latest DOT "
                "physical; envelopes and prices from %s's fit report of 28 August 2026; body dimensions from the %s "
                "(%d), within sex, moved to the operators Halvard measured in August 2026.".format(ops["n"])
                % (W.STANDARD, W.MAKER.split()[0], W.SURVEY, W.SURVEY_YEAR))
    ws["A3"].font = Font(italic=True, color=GREYTXT, size=9)
    ws.append([])
    heads = ["Configuration", "Options", "Price per bus", "Bus operators it fits", "At least 95 per cent",
             "Twelve-year cost per bus"]
    ws.append(heads)
    for c in ws[5]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[5].height = 32
    for c, opts, price in W.CONFIGS:
        ws.append([c, "; ".join(W.OPTIONS[o] for o in opts) or "Standard workstation", price,
                   round(fig["acc_pct"][c] / 100.0, 3), "Yes" if fig["meets"][c] else "No", fig["twelve"][c]])
        row = ws[ws.max_row]
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[2].number_format = "$#,##0"
        row[3].number_format = "0.0%"
        row[4].alignment = Alignment(horizontal="center")
        row[5].number_format = "$#,##0"
        if c == call["config"]:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=PALE)
                cell.font = Font(bold=True)
    r0 = ws.max_row + 2
    lines = [
        ("Bus operators counted", "{:,} on the roster of 30 September 2026, full-time and part-time, every duty status "
         "({:,} men, {:,} women)".format(ops["n"], ops["men"], ops["women"])),
        ("Share fitted", "each operator's probability that all five fit dimensions lie within the envelope, under the "
         "survey's within-sex normal distribution of the five given the operator's stature and weight at the latest "
         "DOT physical, its eye height raised by how far the {:,} operators Halvard measured in August 2026 sit "
         "above the survey's eye height for their own stature and weight ({:+.1f} mm for men, {:+.1f} mm for "
         "women), averaged over the operators".format(
             fig["trial"]["n"], fig["trial"]["eye_height_offset_mm"]["M"], fig["trial"]["eye_height_offset_mm"]["F"])),
        ("Nearest cheaper configuration",
         ("%s, the next price down, %.1f per cent" % (call["next_cheaper"], call["next_cheaper_pct"]))
         if call["nearest_cheaper"] == call["next_cheaper"] else
         ("%s, %.1f per cent; %s, the next price down, %.1f per cent"
          % (call["nearest_cheaper"], call["nearest_cheaper_pct"], call["next_cheaper"], call["next_cheaper_pct"]))),
        ("Twelve-year cost per bus", "the price per bus plus twelve years of each option's maintenance at its "
         "2021 to 2025 cost per bus-year (sheet Option maintenance)"),
    ]
    for k, (lab, val) in enumerate(lines):
        ws.cell(row=r0 + k, column=1, value=lab).font = Font(bold=True)
        ws.cell(row=r0 + k, column=2, value=val).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r0 + k, start_column=2, end_row=r0 + k, end_column=6)
        ws.row_dimensions[r0 + k].height = 30 if len(val) > 110 else 15
    for i, wd in enumerate((24, 46, 13, 14, 12, 15), 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = "A6"
    ws.print_title_rows = "5:5"

    mt = wb.create_sheet("Option maintenance")
    mt["A1"] = "Operator workstation options: maintenance, 2021 to 2025"
    mt["A1"].font = Font(bold=True, size=11)
    mt["A2"] = ("Work orders billed to TVRTA and closed in 2021 to 2025 on each option's component; bus-years are the "
                "months a bus in the fleet carried the option on the first day, divided by twelve.")
    mt["A2"].font = Font(italic=True, color=GREYTXT, size=9)
    mt.append([])
    mt.append(["Option", "Component", "Maintenance cost, 2021 to 2025", "Bus-years fitted", "Cost per bus-year"])
    for c in mt[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    mt.row_dimensions[4].height = 32
    comp = {v[1]: k for k, v in W.COMPONENTS.items() if v[1]}
    for o in W.OPTION_ORDER:
        m = fig["maint"][o]
        mt.append([W.OPTIONS[o], comp[o], round(m["cost"]), m["bus_years"], m["rate"]])
        row = mt[mt.max_row]
        row[2].number_format = "$#,##0"
        row[3].number_format = "#,##0.0"
        row[4].number_format = "$#,##0"
    for i, wd in enumerate((32, 12, 18, 13, 13), 1):
        mt.column_dimensions[get_column_letter(i)].width = wd
    mt.freeze_panes = "A5"

    nt = wb.create_sheet("Notes")
    notes = [
        "Prepared by Fleet Engineering for the Bus Specification Committee's meeting of 29 October 2026, ahead of the "
        "Board's adoption of the FY2027 capital budget on 19 November 2026. The configuration has to be confirmed to "
        "%s by 4 December 2026." % W.MAKER,
        "Standard %s, 3.1 to 3.3: a configuration fits an operator when all five fit dimensions (popliteal height, "
        "buttock-knee length, eye height sitting, thumb-tip reach, abdominal extension depth sitting) lie within its "
        "envelope; the Authority specifies the lowest-priced configuration that fits at least 95 per cent of its bus "
        "operators." % W.STANDARD,
        "The fit dimensions are not measured on our operators, apart from the sitting eye height of the %d Halvard "
        "measured at Ridgeway in August 2026. Their stature and weight are, at each DOT physical (%s to June 2021, %s "
        "since); heights in inches and weights in pounds are taken in millimetres and kilograms."
        % (W.TRIAL_N, W.OLD, W.NEW),
        "Shares are compared with 95 per cent unrounded and shown to one decimal.",
        "Option maintenance: parts and labor on work orders billed to TVRTA, counted once for each work order's "
        "labor, by the month the order closed; work billed to Halvard's warranty or to the 2024 retrofit project is "
        "not maintenance. The buses disposed of since 2021 count for the months they were in the fleet; the pedal "
        "sets fitted in the 2024 retrofit count from the month after they were fitted.",
        "Money is shown to the dollar; the twelve-year figures are the price plus twelve times the sum of the "
        "options' costs per bus-year as shown.",
    ]
    nt["A1"] = "Notes"
    nt["A1"].font = Font(bold=True, size=11)
    for k, t in enumerate(notes, 3):
        c = nt.cell(row=k, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        nt.row_dimensions[k].height = 44
    nt.column_dimensions["A"].width = 110
    wb.save(path)


def write_chart(path, fig, fmt="svg"):
    call = fig["call"]["config"]
    order = [c for c, _, _ in W.CONFIGS][::-1]
    vals = [fig["acc_pct"][c] for c in order]
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9.5, "svg.fonttype": "none",
                         "svg.hashsalt": "task256"})
    f, ax = plt.subplots(figsize=(8.4, 4.9))
    ys = range(len(order))
    for y, c, v in zip(ys, order, vals):
        col = "#" + TEAL if c == call else ("#9AA7B4" if v < 95 else "#6E8FA8")
        ax.plot([75, v], [y, y], color=col, lw=2.2 if c == call else 1.4, zorder=2, solid_capstyle="butt")
        ax.scatter([v], [y], s=64 if c == call else 40, color=col, zorder=3)
        left = v < 95 and v + 2.4 > 95          # a label the standard's line would cross sits left of its dot
        ax.text(v - 0.35 if left else v + 0.35, y + (0.24 if left else 0), "%.1f%%" % v, va="center",
                ha="right" if left else "left", fontsize=8.8, fontweight="bold" if c == call else "normal",
                color="#1F2329")
    ax.axvline(95, color="#B23A30", lw=1.3, ls="--", zorder=1)
    ax.text(95.15, len(order) - 0.45, "Standard ES-12: 95 per cent", color="#B23A30", fontsize=8.4, va="bottom")
    ax.set_yticks(list(ys))
    ax.set_yticklabels(["%s  ${:,}".format(fig["prices"][c]) % c for c in order])
    for lab, c in zip(ax.get_yticklabels(), order):
        if c == call:
            lab.set_fontweight("bold")
            lab.set_color("#" + TEAL)
    ax.set_xlim(75, 100)
    ax.set_ylim(-0.6, len(order) - 0.1)
    ax.set_xlabel("Share of the Authority's bus operators the configuration fits (per cent)")
    ax.set_ylabel("Configuration and price per bus")
    ax.grid(axis="x", color="#E3E6EA", zorder=0)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    y_call = order.index(call)
    ax.annotate("Recommended: %s" % call, xy=(fig["acc_pct"][call] - 0.15, y_call + 0.12),
                xytext=(fig["acc_pct"][call] - 6.2, y_call + 0.62),
                arrowprops=dict(arrowstyle="->", color="#" + TEAL, shrinkA=2, shrinkB=4), color="#" + TEAL,
                fontsize=9.5, fontweight="bold", va="center")
    ax.set_title("Configuration %s is the lowest-priced configuration that fits 95 per cent of our operators" % call,
                 fontsize=10.5, loc="left")
    f.text(0.01, 0.012, "{:,} bus operators on the roster of 30 September 2026 at their latest DOT physical; envelopes "
           "and prices from Halvard's fit report, 28 August 2026;\nbody dimensions from the Brennard survey of 2012, "
           "within sex, given each operator's stature and weight, moved to the operators Halvard measured in August "
           "2026.".format(fig["operators"]["n"]), fontsize=7.2,
           color="#" + GREYTXT)
    f.tight_layout(rect=(0, 0.06, 1, 1))
    if fmt == "svg":
        f.savefig(path, format="svg", metadata={"Date": None, "Creator": None, "Format": None, "Type": None})
    else:
        f.savefig(path, format=fmt, dpi=130)
    plt.close(f)
