"""Independent verifier: reads only the shipped bundle (target/) and the goldens, recomputes the call, the four years'
bridge, each year's file at its close and the quarterly EMS measure with its own code, and checks the goldens state
the same figures."""
import collections
import csv
import datetime as dt
import os
import re
import sys

import numpy as np

UCOD = ("X40", "X41", "X42", "X43", "X44", "X60", "X61", "X62", "X63", "X64", "X85", "Y10", "Y11", "Y12", "Y13", "Y14")
OPI = ("T40.0", "T40.1", "T40.2", "T40.3", "T40.4", "T40.6")
STI = ("T40.5", "T43.6")
OPI_DRUGS = {"fentanyl", "heroin", "oxycodone", "hydrocodone", "morphine", "methadone", "tramadol", "opioid (unspecified)"}
STI_DRUGS = {"methamphetamine", "amphetamine", "cocaine"}
CLOSE = {}


def rd(p):
    with open(p, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def dot(c):
    return c if "." in c or len(c) < 4 else c[:3] + "." + c[3:]


def cat(v):
    cs = [dot(c) for c in v["record_axis"].split()]
    return "OL" if any(c.startswith(OPI) for c in cs) else ("SO" if any(c.startswith(STI) for c in cs) else "N")


def main(target, golden):
    import openpyxl
    deaths = rd(os.path.join(target, "wendmoor_death_records_2019_2026.csv"))
    rx = rd(os.path.join(target, "interstate_exchange_receipts_2019_2027.csv"))
    rec = {(r["state_file_no"], r["record_version"]): r["received_date"] for r in rx}
    cert = {r["state_file_no"]: (r["sending_state"], r["sending_state_file_no"]) for r in rx}
    for r in rd(os.path.join(target, "vital_records_data_year_status.csv")):
        CLOSE[int(r["data_year"])] = dt.date.fromisoformat(r["closed_date"])
    by = collections.defaultdict(list)
    for r in deaths:
        r = dict(r)
        r["v"] = int(r["record_version"])
        r["in"] = dt.date.fromisoformat(rec.get((r["state_file_no"], r["record_version"]), r["registered_date"]))
        by[r["state_file_no"]].append(r)

    def stand(day=None):
        out = []
        for s, vs in by.items():
            vs = sorted([v for v in vs if day is None or v["in"] <= day], key=lambda v: v["v"])
            if vs and vs[-1]["record_status"] != "Void":
                out.append(vs[-1])
        return out

    def od_res(v):
        return v["underlying_cause"][:3] in UCOD and v["residence_county"] == "Wendmoor" and not v["residence_state"]

    now = [v for v in stand() if od_res(v)]
    final = {v["state_file_no"]: v for v in stand()}
    wb = openpyxl.load_workbook(os.path.join(target, "ome_drug_death_review_2025.xlsx"), read_only=True)
    it = wb["Cases"].iter_rows(values_only=True)
    head = next(it)
    review = [dict(zip(head, row)) for row in it if row[0]]
    m = {c: collections.Counter() for c in "OSR"}
    for r in review:
        drugs = {x.strip() for x in (r["drugs_caused_or_contributed"] or "").split(";")}
        c = "O" if drugs & OPI_DRUGS else ("S" if drugs & STI_DRUGS else "R")
        m[c][cat(final[r["state_file_no"]])] += 1
    P = np.array([[m[c][k] / sum(m[c].values()) for k in ("OL", "SO", "N")] for c in "OSR"])
    t = collections.defaultdict(collections.Counter)
    for v in now:
        t[int(v["date_of_death"][:4])][cat(v)] += 1
    est = {y: np.linalg.solve(P.T, np.array([t[y][k] for k in ("OL", "SO", "N")], dtype=float)) for y in t}
    o = {y: est[y][0] for y in est}
    base = np.mean([o[y] for y in (2021, 2022, 2023)])
    ch = 100 * (1 - o[2026] / base)
    r0 = lambda x: int(np.floor(x + .5))
    problems = []
    call = "release" if ch >= 20 else ("first_year" if ch > 0 else "end")
    if call != "release":
        problems.append("call %s" % call)
    bridge = {y: (t[y]["OL"], r0(P[0][2] * o[y]), r0(t[y]["OL"] - P[0][0] * o[y]), r0(o[y])) for y in (2021, 2022, 2023, 2026)}
    # each year's deaths on the file now, among them at the close or joined since; one death per sending certificate
    death = lambda sfn: cert.get(sfn, sfn)
    today = collections.defaultdict(set)
    for v in now:
        today[int(v["date_of_death"][:4])].add(death(v["state_file_no"]))
    ac = {}
    for y in range(2019, 2025):
        atc = {death(v["state_file_no"]) for v in stand(CLOSE[y]) if od_res(v) and int(v["date_of_death"][:4]) == y}
        ac[y] = (len(today[y]), len(today[y] & atc), len(today[y] - atc))
    # EMS encounters: a report handing the patient on joins the report of the unit it names
    ems = rd(os.path.join(target, "ems_patient_care_reports_2021_2026.csv"))
    hand = "Patient Treated, Transferred Care to Another EMS Unit"
    by_inc = collections.defaultdict(list)
    for r in ems:
        by_inc[r["incident_number"]].append(r)

    def poisoning(code):
        c = code.replace(".", "")
        return c[:4] in ("T400", "T401", "T402", "T403", "T404", "T406") and len(c) >= 6 and c[5] in "1234"
    q = collections.Counter()
    for inc, rs in by_inc.items():
        heads = [r for r in rs if r["disposition"] != hand]
        groups = {id(h): [h] for h in heads}
        for r in rs:
            if r["disposition"] == hand:
                hs = [h for h in heads if "%s %s" % (h["agency_name"], h["unit"]) == r["destination"]]
                if len(hs) != 1:
                    problems.append("EMS %s: %s names %d reports" % (inc, r["pcr_number"], len(hs)))
                    continue
                groups[id(hs[0])].append(r)
        for g in groups.values():
            crit = any(r["naloxone_administered"] == "Yes" or poisoning(r["primary_impression"]) or
                       poisoning(r["secondary_impression"]) for r in g)
            res = any(r["patient_home_county"] == "Wendmoor" and not r["patient_home_state"] for r in g)
            alive = not any(r["disposition"].startswith("Dead at Scene") for r in g)
            if crit and res and alive:
                d = g[0]["psap_call_datetime"]
                q["%sQ%d" % (d[:4], (int(d[5:7]) - 1) // 3 + 1)] += 1
    # the goldens state the same figures
    gw = openpyxl.load_workbook(os.path.join(golden, "ori_outcome_measures_2019_2026.xlsx"), data_only=True)
    ws = gw["File at close"]
    seen = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if isinstance(row[0], int):
            seen.add(row[0])
            if tuple(row[2:5]) != ac[row[0]]:
                problems.append("workbook %d file now %s, recomputed %s" % (row[0], row[2:5], ac[row[0]]))
    if seen != set(ac):
        problems.append("workbook years %s" % sorted(seen))
    es = gw["EMS non-fatal"]
    got = {row[0]: row[1] for row in es.iter_rows(min_row=2, values_only=True) if row[0] and re.match(r"\d{4}Q\d", str(row[0]))}
    if got != dict(q):
        problems.append("workbook EMS differs in %d quarters" % sum(1 for k in q if got.get(k) != q[k]))
    from pypdf import PdfReader
    text = " ".join(p.extract_text() for p in PdfReader(os.path.join(golden, "ori_phase_one_outcome_certification_2026.pdf")).pages)
    text = re.sub(r"\s+", " ", text)
    for needle in ("%.1f per cent below the baseline" % ch, "certifies that %d Wendmoor County residents" % r0(o[2026]),
                   "baseline of %d a year" % r0(base)):
        if needle not in text:
            problems.append("certification lacks %r" % needle)
    for y, b in bridge.items():
        if not re.search(r"%d %s %s %s %s" % (y, *["{:,}".format(x) for x in b]), text):
            problems.append("certification bridge %d %s not found" % (y, b))
    print("verify: change %.3f (%s), 2026 %.2f, baseline %.2f; bridge %s; at close %s; EMS %d quarters; %d problems" % (
        ch, call, o[2026], base, bridge, ac, len(q), len(problems)))
    for p in problems:
        print("  PROBLEM", p)
    return not problems


if __name__ == "__main__":
    t = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "target")
    g = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(t), "golden")
    sys.exit(0 if main(t, g) else 1)
