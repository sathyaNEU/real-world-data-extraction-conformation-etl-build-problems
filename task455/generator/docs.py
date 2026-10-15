"""The pack's documents: the Department of Commerce's designation guideline for the Formation Counties Program,
the Labor Market Information Bureau's data dictionary for the research extract, and the extract manifest."""
import csv

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import letter                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import mm                                                    # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import world as W   # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("LiberationSans", "Regular"), ("LiberationSans-Bold", "Bold"), ("LiberationSans-Italic", "Italic")):
    pdfmetrics.registerFont(TTFont(nm, FONT_DIR + nm.split("-")[0] + "-" + f + ".ttf"))
INK = colors.HexColor("#1d2329")
GREY = colors.HexColor("#5a636b")
NAVY = colors.HexColor("#1f3a5c")
BODY = ParagraphStyle("b", fontName="LiberationSans", fontSize=10, leading=13.8, textColor=INK, spaceAfter=6)
NUM = ParagraphStyle("n", parent=BODY, leftIndent=28, firstLineIndent=-28)
TITLE = ParagraphStyle("t", fontName="LiberationSans-Bold", fontSize=14, leading=18, textColor=NAVY, spaceAfter=2)
SUB = ParagraphStyle("s", fontName="LiberationSans", fontSize=10, leading=13.5, textColor=GREY, spaceAfter=12)
H2 = ParagraphStyle("h2", fontName="LiberationSans-Bold", fontSize=10.8, leading=14, textColor=NAVY, spaceBefore=7,
                    spaceAfter=3)
CELL = ParagraphStyle("c", fontName="LiberationSans", fontSize=9.4, leading=12.2, textColor=INK)
REF = "FC-G-2027-01"


def write_guideline(path):
    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFont("LiberationSans", 8)
        canvas.setFillColor(GREY)
        canvas.drawString(22 * mm, 12 * mm, "Formation Counties Program, Designation Guideline %s" % REF)
        canvas.drawRightString(letter[0] - 22 * mm, 12 * mm, "Page %d of 2" % doc.page)
        canvas.restoreState()
    d = SimpleDocTemplate(path, pagesize=letter, leftMargin=22 * mm, rightMargin=22 * mm, topMargin=20 * mm,
                          bottomMargin=20 * mm, title="Formation Counties Program Designation Guideline %s" % REF,
                          author=W.ORG_COMMERCE, subject="", creator="", producer="")
    P = lambda t: Paragraph(t, NUM)
    s = [Paragraph("Department of Commerce &nbsp;|&nbsp; Office of Business Development", SUB),
         Paragraph("Formation Counties Program", TITLE),
         Paragraph("Designation Guideline for the 2027 to 2029 Grant Period (%s), issued 2 February 2027" % REF, SUB)]
    s.append(Paragraph("1. Purpose", H2))
    s.append(P("1.1&nbsp;&nbsp;The Formation Counties Program helps the counties where new businesses start to keep "
               "them: it funds county development authorities' small business advisers, microloan funds and "
               "incubator space."))
    s.append(P("1.2&nbsp;&nbsp;This guideline sets how the Secretary of Commerce designates the Formation Counties of the "
               "Lornock Valley Development Region for the grant period 1 July 2027 to 30 June 2029, under the "
               "appropriation the legislature made for the program in its 2027 session."))
    s.append(Paragraph("2. Designation", H2))
    s.append(P("2.1&nbsp;&nbsp;The Secretary designates three of the region's seven counties (Brennock, Caldmere, "
               "Dunrath, Eskwell, Fennamore, Quarrin and Tolbeck): the three with the highest establishment birth "
               "rates in calendar year 2026, each county's rate measured from its own establishments."))
    s.append(P("2.2&nbsp;&nbsp;The designation is made by 30 June 2027. Each designated county's development authority "
               "receives a grant of $1,250,000 for the grant period."))
    s.append(P("2.3&nbsp;&nbsp;If a designated county's authority declines the grant or cannot meet its conditions by 1 "
               "September 2027, the county with the next highest rate is designated in its place."))
    s.append(Paragraph("3. Definitions", H2))
    s.append(P("3.1&nbsp;&nbsp;<b>Establishment.</b> A single worksite at which a private employer has employment "
               "covered by the state's unemployment insurance law."))
    s.append(P("3.2&nbsp;&nbsp;<b>County.</b> An establishment is counted in the county in which its worksite lies."))
    s.append(P("3.3&nbsp;&nbsp;<b>Birth.</b> An establishment that first reported employment in 2026 and whose business "
               "did not operate before under another employer account. A business that is sold, reorganised or "
               "changes its legal form and goes on under a new employer account continues; it is not a birth."))
    s.append(P("3.4&nbsp;&nbsp;<b>Birth rate.</b> A county's births per 100 of its private establishments that reported "
               "employment in at least one month of the fourth quarter of 2025."))
    s.append(Paragraph("4. Rates", H2))
    s.append(P("4.1&nbsp;&nbsp;Rates are computed from the unemployment insurance program's records as the Department "
               "of Workforce Services' Labor Market Information Bureau holds them after the reports for the fourth "
               "quarter of 2026 are due on 30 April 2027."))
    s.append(P("4.2&nbsp;&nbsp;Rates are compared unrounded. The designation memorandum shows them to one decimal."))
    s.append(Paragraph("5. Data and confidentiality", H2))
    s.append(P("5.1&nbsp;&nbsp;The Labor Market Information Bureau prepares the rates under the data sharing agreement "
               "between the Department of Commerce and the Department of Workforce Services (DWS-DOC-2023-07). Records "
               "that identify an employer or a person stay with the bureau; the memorandum to the Secretary reports "
               "county figures."))
    s.append(P("5.2&nbsp;&nbsp;Grant conditions, reporting and payment schedules are set in the grant agreements and are "
               "not part of this guideline."))
    s.append(Spacer(1, 10))
    sig = Table([[Paragraph("Issued for the Secretary of Commerce<br/><br/>%s, Program Director<br/>Office of Business "
                            "Development<br/>2 February 2027" % W.PEOPLE["director"], CELL)]], colWidths=[90 * mm],
                hAlign="LEFT")
    sig.setStyle(TableStyle([("LINEABOVE", (0, 0), (-1, 0), 0.5, GREY), ("TOPPADDING", (0, 0), (-1, -1), 6),
                             ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    s.append(sig)
    s.append(Paragraph("Annex A. Timetable for the 2027 to 2029 grant period", H2))
    rows = [("2 February 2027", "Guideline issued"),
            ("30 April 2027", "Reports for the fourth quarter of 2026 due from employers"),
            ("May 2027", "Labor Market Information Bureau computes the 2026 rates"),
            ("by 30 June 2027", "Secretary designates the Formation Counties"),
            ("1 July 2027", "Grant period begins"),
            ("1 September 2027", "Grant agreements signed; a county that declines is replaced under 2.3"),
            ("30 June 2029", "Grant period ends")]
    tt = Table([[Paragraph("<b>Date</b>", CELL), Paragraph("<b>Step</b>", CELL)]] +
               [[Paragraph(a, CELL), Paragraph(b, CELL)] for a, b in rows], colWidths=[38 * mm, 120 * mm],
               hAlign="LEFT")
    tt.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, 0), 0.5, NAVY), ("TOPPADDING", (0, 0), (-1, -1), 2),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 2), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    s.append(tt)
    s.append(Paragraph("Annex B. County development authorities eligible to receive a grant", H2))
    auth = [("Brennock", "Brennock County Economic Development Authority"),
            ("Caldmere", "Caldmere County Development Commission"),
            ("Dunrath", "Dunrath Area Economic Development Council"),
            ("Eskwell", "Eskwell Area Development Partnership"),
            ("Fennamore", "Fennamore County Development Corporation"),
            ("Quarrin", "Quarrin County Economic Development Authority"),
            ("Tolbeck", "Tolbeck County Industrial Development Authority")]
    ta = Table([[Paragraph("<b>County</b>", CELL), Paragraph("<b>Authority</b>", CELL)]] +
               [[Paragraph(a, CELL), Paragraph(b, CELL)] for a, b in auth], colWidths=[38 * mm, 120 * mm],
               hAlign="LEFT")
    ta.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, 0), 0.5, NAVY), ("TOPPADDING", (0, 0), (-1, -1), 2),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 2), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    s.append(ta)
    d.build(s, onFirstPage=frame, onLaterPages=frame)


FILE_ROWS = {
    "register": "Every private employer account with a worksite in the seven counties that was open at any time from 1 "
                "October 2024 to the extract, including accounts opened that have not yet reported employment. One row "
                "per account, as the register stood on 14 May 2027.",
    "qcew": "The quarterly contribution reports (QCEW) of those accounts, fourth quarter of 2024 to fourth quarter of "
            "2026. One row per account and quarter in which the account was open; an open account reports every "
            "quarter, with zero employment when it employed no one. Reports for the fourth quarter of 2026 as received "
            "by the extract.",
    "wages": "The quarterly wage reports of those accounts: one row per worker, account and quarter in which the account "
             "paid the worker wages. Pipe-delimited text with a header line.",
    "counties": "The region's county codes.",
    "guideline": "Department of Commerce guideline FC-G-2027-01 (designation of Formation Counties for 2027 to 2029).",
    "dict": "This dictionary.",
    "manifest": "The list of the files in the extract.",
    "pop": "The state demographer's July 1 population estimates for the seven counties, vintage 2026.",
    "sos": "New business entity filings with the Secretary of State, 2025 and 2026, with a principal office in the "
           "seven counties.",
    "grants": "Formation Counties Program grant disbursements in fiscal year 2026 (the 2025 to 2027 grant period).",
}

FIELDS = [
    ("register", "account_no", "UI employer account number. Accounts opened since the UI tax system replacement went live "
                               "on 1 January 2025 are numbered from 800000000 in the order they were opened."),
    ("register", "legal_name", "Legal name of the employing entity."),
    ("register", "trade_name", "Name the employer does business under at the worksite."),
    ("register", "fein", "Federal employer identification number."),
    ("register", "worksite_address", "Street address of the worksite. Accounts carried over from the previous system "
                                     "keep the address as it was keyed there; accounts opened since 1 January 2025 carry "
                                     "it as the new system standardises it."),
    ("register", "worksite_city", "Postal city of the worksite."),
    ("register", "worksite_zip", "ZIP code of the worksite."),
    ("register", "county_code", "County in which the worksite lies, geocoded from the worksite address (see Codes)."),
    ("register", "naics_code", "Industry code (NAICS) on the account."),
    ("register", "naics_vintage", "NAICS edition of naics_code: 2017 for accounts carried over from the previous system, "
                                  "2022 for accounts opened since."),
    ("register", "ownership_code", "Ownership: 5 private."),
    ("register", "registration_date", "Date the account was opened."),
    ("register", "liability_date", "Date of the first payroll that made the employer liable under the state's "
                                   "unemployment insurance law."),
    ("register", "status", "ACTIVE or INACTIVE on the extract date."),
    ("register", "inactivation_date", "Date the account was closed. An account stays open until the employer reports "
                                      "that it has stopped employing or the tax unit closes it after four quarters "
                                      "without employment."),
    ("register", "predecessor_account", "Account whose experience rating was transferred to this account on the "
                                        "employer's application (Form UI-22), recorded when the tax unit approves the "
                                        "transfer. A transfer is required where the businesses are under common "
                                        "ownership, management or control and is otherwise optional; an employer that "
                                        "takes over a business without applying has no predecessor recorded."),
    ("register", "transfer_date", "Effective date of the experience transfer."),
    ("qcew", "account_no", "UI employer account number."),
    ("qcew", "year", "Year of the report."),
    ("qcew", "qtr", "Quarter of the report (1 to 4)."),
    ("qcew", "county_code", "County of the worksite as assigned for the quarter: from the worksite ZIP code through the "
                            "second quarter of 2025, and by geocoding the worksite street address from the third quarter "
                            "of 2025."),
    ("qcew", "naics_code", "Industry code for the quarter (NAICS 2022)."),
    ("qcew", "own_code", "Ownership: 5 private."),
    ("qcew", "month1_emplvl", "Workers on the payroll for the pay period including the 12th of the first month of the "
                              "quarter."),
    ("qcew", "month2_emplvl", "The same for the second month."),
    ("qcew", "month3_emplvl", "The same for the third month."),
    ("qcew", "total_qtrly_wages", "Wages paid in the quarter, dollars."),
    ("qcew", "taxable_qtrly_wages", "Wages paid in the quarter up to the taxable wage base of $14,500 a worker a year, "
                                    "dollars."),
    ("wages", "worker_key", "The bureau's protected key for the worker's Social Security number, the same for a worker "
                            "in every account and quarter."),
    ("wages", "account_no", "UI employer account number of the paying account."),
    ("wages", "year", "Year the wages were paid."),
    ("wages", "qtr", "Quarter the wages were paid (1 to 4)."),
    ("wages", "wages", "Wages the account paid the worker in the quarter, dollars."),
    ("counties", "county_code", "Three-digit county code used in the register and the reports."),
    ("counties", "county_name", "County name."),
    ("counties", "region", "Development region."),
    ("counties", "county_seat", "County seat."),
]


def write_dictionary(path, counts):
    import openpyxl
    from openpyxl.styles import Alignment, Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Files"
    ws.append(["Lornock Valley research extract of 14 May 2027: files"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Labor Market Information Bureau, Business Dynamics Unit. Prepared for the Formation Counties "
               "designation under data sharing agreement DWS-DOC-2023-07."])
    ws.append([])
    ws.append(["File", "Contents", "Rows"])
    for c in ws[4]:
        c.font = Font(bold=True)
    for k in ("register", "qcew", "wages", "counties", "guideline", "pop", "sos", "grants", "manifest", "dict"):
        ws.append([W.FILES[k], FILE_ROWS[k], counts.get(k)])
    for row in ws.iter_rows(min_row=5):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[0].alignment = Alignment(vertical="top")
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 96
    ws.column_dimensions["C"].width = 10
    f = wb.create_sheet("Fields")
    f.append(["File", "Field", "Description"])
    for c in f[1]:
        c.font = Font(bold=True)
    for k, fld, desc in FIELDS:
        f.append([W.FILES[k], fld, desc])
        f.cell(f.max_row, 3).alignment = Alignment(wrap_text=True, vertical="top")
    f.column_dimensions["A"].width = 46
    f.column_dimensions["B"].width = 22
    f.column_dimensions["C"].width = 100
    c = wb.create_sheet("Codes")
    c.append(["Field", "Code", "Meaning"])
    for cell in c[1]:
        cell.font = Font(bold=True)
    for code, name, seat, _ in W.COUNTIES:
        c.append(["county_code", code, "%s County" % name])
    c.append(["ownership_code / own_code", "5", "Private"])
    c.append(["status", "ACTIVE", "Account open on the extract date"])
    c.append(["status", "INACTIVE", "Account closed (inactivation_date)"])
    c.append(["naics_vintage", "2017", "North American Industry Classification System, 2017 edition"])
    c.append(["naics_vintage", "2022", "North American Industry Classification System, 2022 edition"])
    c.column_dimensions["A"].width = 26
    c.column_dimensions["B"].width = 10
    c.column_dimensions["C"].width = 60
    wb.save(path)


def write_manifest(path, counts):
    rows = [("register", "UI employer register", "accounts open 2024-10-01 to 2027-05-14", "UI tax system (employer accounts)"),
            ("qcew", "Quarterly contribution reports", "2024Q4 to 2026Q4", "UI tax system (contribution reports)"),
            ("wages", "Quarterly wage reports", "2024Q4 to 2026Q4", "UI wage record system"),
            ("counties", "County codes", "", "LMI Bureau reference tables"),
            ("guideline", "Program guideline FC-G-2027-01", "issued 2027-02-02", "Department of Commerce"),
            ("pop", "County population estimates", "2020 to 2026", "State demographer"),
            ("sos", "Secretary of State entity filings", "2025-01-01 to 2026-12-31", "Secretary of State business filings"),
            ("grants", "Formation Counties grant disbursements", "FY2026", "Department of Commerce, Office of Business Development"),
            ("dict", "Data dictionary", "", "LMI Bureau"),
            ]
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\r\n")
        w.writerow(["file_name", "title", "rows", "period", "source", "prepared_by", "extract_date"])
        for k, title, period, src in rows:
            w.writerow([W.FILES[k], title, counts.get(k, ""), period, src,
                        "LMI Bureau, Business Dynamics Unit" if k in ("register", "qcew", "wages", "counties", "dict")
                        else "received by the LMI Bureau", W.EXTRACT.isoformat()])
        w.writerow([W.FILES["manifest"], "Extract manifest", len(rows) + 1, "", "LMI Bureau",
                    "LMI Bureau, Business Dynamics Unit", W.EXTRACT.isoformat()])
    return len(rows) + 1
