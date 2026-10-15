"""The goldens: the department's phase one outcome certification (PDF) and the outcome measures workbook, both from
the figures computed here off the pack, and generator/figures.json with every figure the write-up states."""
import collections
import datetime as dt
import json
import os

import numpy as np

import conform as C
import methods as M
import world as W

CERT = "ori_phase_one_outcome_certification_2026.pdf"
WORKBOOK = "ori_outcome_measures_2019_2026.xlsx"


WORDS = {k: w for k, w in enumerate(["No", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine", "Ten", "Eleven", "Twelve"])}


def rnd(x):
    return int(np.floor(x + 0.5))


# ---------------------------------------------------------------- Ask A: each year's file at its close

ASK_A_YEARS = list(range(2019, 2025))


def after_close(recs, receipts):
    """For each year, the residents' drug overdose deaths on the file now, split into those already among them on
    the file the day it closed and those that joined since; a death counted once, by its sending state's certificate
    number when it came through the exchange."""
    key = C.death_keys(receipts)
    k = lambda v: key.get(v["state_file_no"], ("S", v["state_file_no"]))
    today = collections.defaultdict(set)
    for v in C.resident_overdoses(recs):
        today[C.year_of(v)].add(k(v))
    out = {}
    for y in ASK_A_YEARS:
        atc = {k(v) for v in C.resident_overdoses(recs, W.CLOSE[y]) if C.year_of(v) == y}
        out[y] = {"today": len(today[y]), "among": len(today[y] & atc), "joined": len(today[y] - atc)}
    return out


# ---------------------------------------------------------------- Ask B: EMS non-fatal opioid overdose encounters

def ems_rows(target):
    from pack import FILES
    return C.read_csv(os.path.join(target, FILES["ems"]))


def ems_encounters(rows):
    import ems as EM
    return EM.count(rows)


# ---------------------------------------------------------------- the figures

def figures(target):
    L = C.load(target)
    recs, review = L["recs"], L["review"]
    g, est, P, t = M.golden(recs, review)
    pending = sum(1 for v in C.as_of(recs) if C.is_resident(v) and C.year_of(v) == W.TARGET
                  and v["manner_of_death"] == "Pending investigation")
    m = M.review_matrix(recs, review)[None]
    base = np.mean([g[y] for y in W.BASE])
    ch = 100 * (1 - g[W.TARGET] / base)
    call = "release" if ch >= W.THRESHOLD else ("first_year" if ch > 0 else "end")
    bridge = {}
    for y in W.BASE + [W.TARGET]:
        o = g[y]
        bridge[y] = {"certified": t[y]["OL"], "missed": P[0][2] * o, "false": t[y]["OL"] - P[0][0] * o, "outcome": o}
    f = {
        "call": call,
        "change_pct": round(ch, 1),
        "outcome_2026": rnd(g[W.TARGET]),
        "baseline": rnd(base),
        "baseline_years": {str(y): rnd(g[y]) for y in W.BASE},
        "bridge": {str(y): {k: (v if k == "certified" else rnd(v)) for k, v in b.items()} for y, b in bridge.items()},
        "review": {"deaths": len(review), "opioid": sum(m["O"].values()), "stimulant": sum(m["S"].values()),
                   "other": sum(m["R"].values()), "opioid_certified_with_opioid": m["O"]["OL"],
                   "opioid_certified_no_drug_or_other": m["O"]["N"] + m["O"]["SO"],
                   "stimulant_certified_with_opioid": m["S"]["OL"], "stimulant_certified_stimulant_only": m["S"]["SO"],
                   "stimulant_certified_neither": m["S"]["N"], "other_certified_with_opioid": m["R"]["OL"],
                   "sensitivity": round(P[0][0], 3), "stimulant_listed_with_opioid": round(P[1][0], 3)},
        "certified_baseline_avg": round(np.mean([t[y]["OL"] for y in W.BASE]), 1),
        "certified_2026": t[W.TARGET]["OL"],
        "certified_change_pct": round(100 * (1 - t[W.TARGET]["OL"] / np.mean([t[y]["OL"] for y in W.BASE])), 1),
        "release_line": round(0.8 * base, 1),
        "threshold_pct": W.THRESHOLD,
        "rogan_gladen_change_pct": round(M.change(M.two_by_two(recs, review)[0])[0], 1),
        "rogan_gladen_false_listing_rate": round(M.two_by_two(recs, review)[2], 3),
        "pending_2026": pending,
        "pending_2026_words": WORDS[pending],
        "after_close": {str(y): v for y, v in after_close(recs, L["receipts"]).items()},
        "ems": ems_encounters(ems_rows(target)),
    }
    exact = {"outcome": {str(y): g[y] for y in W.YEARS}, "baseline": base, "change": ch,
             "bridge": {str(y): b for y, b in bridge.items()}, "P": P.tolist(), "counts": {str(y): t[y] for y in t},
             "est": {str(y): est[y].tolist() for y in est}}
    return f, exact


# ---------------------------------------------------------------- the certification PDF

def write_certification(f, ex, path):
    from reportlab import rl_config
    rl_config.invariant = 1
    from reportlab.graphics.shapes import Drawing, Line, Rect, String
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    fd = "/usr/share/fonts/truetype/liberation/"
    for nm, ff in (("LiberationSans", "Regular"), ("LiberationSans-Bold", "Bold"), ("LiberationSans-Italic", "Italic")):
        pdfmetrics.registerFont(TTFont(nm, fd + nm.split("-")[0] + "-" + ff + ".ttf"))
    INK, GREY, NAVY = colors.HexColor("#1e2328"), colors.HexColor("#5d6670"), colors.HexColor("#21456b")
    LIGHT = colors.HexColor("#b9c4cf")
    body = ParagraphStyle("b", fontName="LiberationSans", fontSize=9.8, leading=13.6, textColor=INK, spaceAfter=7)
    small = ParagraphStyle("s", parent=body, fontSize=8.2, leading=11, textColor=GREY, spaceAfter=3)
    head = ParagraphStyle("h", fontName="LiberationSans-Bold", fontSize=14, leading=18, textColor=NAVY, spaceAfter=8)
    h2 = ParagraphStyle("h2", fontName="LiberationSans-Bold", fontSize=10.4, leading=14, textColor=NAVY, spaceBefore=8, spaceAfter=4)
    cell = ParagraphStyle("c", fontName="LiberationSans", fontSize=8.8, leading=11.4, textColor=INK)
    cellb = ParagraphStyle("cb", parent=cell, fontName="LiberationSans-Bold")
    cellr = ParagraphStyle("cr", parent=cell, alignment=2)
    cellrb = ParagraphStyle("crb", parent=cellr, fontName="LiberationSans-Bold")

    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFont("LiberationSans", 7.6)
        canvas.setFillColor(GREY)
        canvas.drawString(20 * mm, 11 * mm, "%s. Phase one outcome certification, 2 June 2027." % W.ORG)
        canvas.drawRightString(letter[0] - 20 * mm, 11 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    d = SimpleDocTemplate(path, pagesize=letter, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=17 * mm, bottomMargin=19 * mm,
                          title="Overdose Reduction Initiative: phase one outcome certification", author=W.ORG, subject="",
                          creator="", producer="")
    s = []
    meta = Table([[Paragraph("<b>To</b>", cell), Paragraph("Board of Trustees, %s (attention: %s, %s)" % (W.FOUNDATION, W.PEOPLE["grants"], W.TITLES["grants"]), cell)],
                  [Paragraph("<b>From</b>", cell), Paragraph("%s, %s, %s" % (W.PEOPLE["manager"], W.TITLES["manager"], W.ORG), cell)],
                  [Paragraph("<b>Date</b>", cell), Paragraph("2 June 2027", cell)],
                  [Paragraph("<b>Subject</b>", cell), Paragraph("Certification of the phase one outcome under Schedule B4, Grant Agreement ORI-2024-01", cell)]],
                 colWidths=[22 * mm, 150 * mm])
    meta.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                              ("TOPPADDING", (0, 0), (-1, -1), 1), ("LINEBELOW", (0, -1), (-1, -1), 0.6, NAVY)]))
    s.append(meta)
    s.append(Spacer(1, 9))
    b = f["bridge"]
    s.append(Paragraph("Residents' opioid overdose deaths in 2026 were %.1f per cent below the 2021 to 2023 baseline: "
                       "Schedule B5 releases phase two" % f["change_pct"], head))
    s.append(Paragraph(
        "Under Schedule B4 the Department certifies that %d Wendmoor County residents died in 2026 of a drug overdose that an "
        "opioid caused or contributed to, against a baseline of %d a year over 2021 to 2023 (%d, %d and %d). The 2026 outcome "
        "is %.1f per cent below the baseline. The fall is at least 20 per cent, so under Schedule B5 the Foundation releases "
        "phase two of the %s." % (f["outcome_2026"], f["baseline"], f["baseline_years"]["2021"], f["baseline_years"]["2022"],
                                   f["baseline_years"]["2023"], f["change_pct"], W.INITIATIVE), body))
    s.append(Paragraph(
        "The certificates alone show far less: %d deaths certified with an opioid in 2026 against %.1f a year in the baseline. "
        "The difference is the certificates' classification error, which Schedule B3 requires the Department to correct in "
        "every year it reports, and which the State Office of the Medical Examiner's review of the 2025 deaths measured." %
        (f["certified_2026"], f["certified_baseline_avg"]), body))
    s.append(Paragraph("Table 1. From the deaths certified with an opioid to the outcome", h2))
    rows = [[Paragraph(x, cellb if i == 0 else cellrb) for i, x in enumerate(
        ["Year", "Certified with an opioid", "Add: opioid deaths certified without one", "Less: certified with an opioid no opioid caused", "Outcome"])]]
    for y in ("2021", "2022", "2023", "2026"):
        rows.append([Paragraph(y, cell)] + [Paragraph("{:,}".format(b[y][k]), cellr) for k in ("certified", "missed", "false")] +
                    [Paragraph("<b>{:,}</b>".format(b[y]["outcome"]), cellr)])
    rows.append([Paragraph("2021 to 2023 average", cell), Paragraph("", cellr), Paragraph("", cellr), Paragraph("", cellr),
                 Paragraph("<b>{:,}</b>".format(f["baseline"]), cellr)])
    tb = Table(rows, colWidths=[40 * mm, 28 * mm, 38 * mm, 42 * mm, 24 * mm])
    tb.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, 0), 0.7, NAVY), ("LINEABOVE", (0, -1), (-1, -1), 0.4, GREY),
                            ("VALIGN", (0, 0), (-1, -1), "BOTTOM"), ("TOPPADDING", (0, 0), (-1, -1), 2.5),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5), ("BACKGROUND", (0, 4), (-1, 4), colors.HexColor("#eef2f6"))]))
    s.append(tb)
    s.append(Paragraph("Deaths of Wendmoor County residents whose underlying cause is a drug poisoning, on the state death file "
                       "as extracted on 11 May 2027. Estimated counts are rounded to the whole death; the change is computed "
                       "from the unrounded counts.", small))
    # the chart: certified against outcome for the four years, with the line the 2026 outcome had to reach
    dw = Drawing(170 * mm, 62 * mm)
    x0, y0, hgt = 18 * mm, 10 * mm, 44 * mm
    top = 300.0
    years = ["2021", "2022", "2023", "2026"]
    for k in (0, 100, 200, 300):
        yy = y0 + hgt * k / top
        dw.add(Line(x0, yy, x0 + 140 * mm, yy, strokeColor=colors.HexColor("#e3e7eb"), strokeWidth=0.5))
        dw.add(String(x0 - 3 * mm, yy - 2.5, str(k), fontName="LiberationSans", fontSize=7, fillColor=GREY, textAnchor="end"))
    for i, y in enumerate(years):
        cx = x0 + (12 + i * 34) * mm
        for j, (key, col) in enumerate((("certified", LIGHT), ("outcome", NAVY))):
            v = b[y][key]
            dw.add(Rect(cx + j * 11 * mm, y0, 10 * mm, hgt * v / top, fillColor=col, strokeColor=None))
            dw.add(String(cx + j * 11 * mm + 5 * mm, y0 + hgt * v / top - 9, "{:,}".format(v), fontName="LiberationSans-Bold",
                          fontSize=7, fillColor=colors.white if key == "outcome" else INK, textAnchor="middle"))
        dw.add(String(cx + 10.5 * mm, y0 - 8, y, fontName="LiberationSans", fontSize=7.6, fillColor=INK, textAnchor="middle"))
    line = 0.8 * ex["baseline"]
    yy = y0 + hgt * line / top
    dw.add(Line(x0, yy, x0 + 140 * mm, yy, strokeColor=colors.HexColor("#b03a2e"), strokeWidth=0.8, strokeDashArray=[3, 2]))
    dw.add(Line(x0 + 92 * mm, y0 + hgt + 4.5 * mm, x0 + 98 * mm, y0 + hgt + 4.5 * mm, strokeColor=colors.HexColor("#b03a2e"),
                strokeWidth=0.8, strokeDashArray=[3, 2]))
    dw.add(String(x0 + 99.5 * mm, y0 + hgt + 3.4 * mm, "20 per cent below the baseline (%.1f)" % f["release_line"],
                  fontName="LiberationSans", fontSize=7, fillColor=INK))
    dw.add(Rect(x0 + 2 * mm, y0 + hgt + 3 * mm, 3 * mm, 3 * mm, fillColor=LIGHT, strokeColor=None))
    dw.add(String(x0 + 6.5 * mm, y0 + hgt + 3.4 * mm, "Certified with an opioid", fontName="LiberationSans", fontSize=7, fillColor=INK))
    dw.add(Rect(x0 + 42 * mm, y0 + hgt + 3 * mm, 3 * mm, 3 * mm, fillColor=NAVY, strokeColor=None))
    dw.add(String(x0 + 46.5 * mm, y0 + hgt + 3.4 * mm, "Outcome (Schedule B1)", fontName="LiberationSans", fontSize=7, fillColor=INK))
    s.append(Paragraph("Figure 1. The certificates fell %.1f per cent; the deaths opioids caused fell %.1f per cent" % (f["certified_change_pct"], f["change_pct"]), h2))
    s.append(dw)
    r = f["review"]
    s.append(Paragraph("How the certificates were corrected", h2))
    s.append(Paragraph(
        "The review examined all %d resident drug overdose deaths on the closed 2025 file. An opioid caused or contributed to "
        "%d of them; the coroner's certificates named an opioid on %d of these and named no drug on the other %d. "
        "Methamphetamine or cocaine caused %d deaths without an opioid; on %d of them the coroner listed fentanyl that "
        "toxicology found but the review found did not contribute, %d were certified with the stimulant alone and %d named "
        "no drug. None of the %d deaths caused by other drugs was certified with an opioid." %
        (r["deaths"], r["opioid"], r["opioid_certified_with_opioid"], r["opioid_certified_no_drug_or_other"], r["stimulant"],
         r["stimulant_certified_with_opioid"], r["stimulant_certified_stimulant_only"], r["stimulant_certified_neither"],
         r["other"]), body))
    s.append(Paragraph(
        "The false listings therefore come from the stimulant deaths, half of which carry an opioid on the certificate, and not "
        "from the other drug deaths. The stimulant deaths have grown every year since 2019 while the other drug deaths have "
        "fallen, so a single false-listing rate taken over all non-opioid deaths would remove too many false listings from the "
        "baseline years and too few from 2026. The Department instead applied the review's rates for each kind of death to "
        "each year's three certificate groups (an opioid named; a stimulant named without an opioid; neither) and solved for "
        "the opioid, stimulant and other drug deaths behind them. The rates are the same in every age group and for both "
        "sexes, so solving within those groups gives the same counts.", body))
    s.append(Paragraph("Notes", h2))
    for t in ("Residents are counted by county of residence on the final record, wherever they died; drug overdose deaths "
              "by an underlying cause of X40 to X44, X60 to X64, X85 or Y10 to Y14 (Schedule B1).",
              "%s 2026 deaths of residents were still pending investigation at the extract; Schedule B4 certifies the file as it stood that day." % f["pending_2026_words"],
              "The secondary measure (Schedule B6) and each year's file at its close are in %s." % WORKBOOK):
        s.append(Paragraph(t, small))
    d.build(s, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------- the workbook

def write_workbook(f, path):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "File at close"
    hdr = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="21456B")
    ws.append(["Data year", "File closed", "Residents' drug overdose deaths on the file at the extract",
               "Among them on the file the day it closed", "Joined since the close"])
    for y in ASK_A_YEARS:
        a = f["after_close"][str(y)]
        ws.append([y, dt.date.fromisoformat(W.CLOSE[y].isoformat()), a["today"], a["among"], a["joined"]])
    for c in ws[1]:
        c.font, c.fill = hdr, fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 58
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "d mmm yyyy"
        for c in row[2:5]:
            c.number_format = "#,##0"
    for col, wdt in zip("ABCDE", (11, 14, 24, 22, 16)):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "A2"
    ws.append([])
    for t in ("The file on a close date is each record's latest version received by the end of that day (an out-of-state "
              "death on the day the interstate exchange delivered it), voids applied.",
              "A death counts once. An out-of-state death is the sending state's certificate: when the sending state sent a "
              "death again as a new record and the office voided the first copy as a duplicate, the death was already on the "
              "file at the close if the first copy was."):
        ws.append([t])
        ws.cell(ws.max_row, 1).font = Font(italic=True, size=9, color="5D6670")
    e = wb.create_sheet("EMS non-fatal")
    e.append(["Quarter", "Non-fatal opioid overdoses of residents treated by EMS"])
    for q, n in f["ems"].items():
        e.append([q, n])
    for c in e[1]:
        c.font, c.fill = hdr, fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    e.row_dimensions[1].height = 32
    for row in e.iter_rows(min_row=2):
        row[1].number_format = "#,##0"
    e.column_dimensions["A"].width = 10
    e.column_dimensions["B"].width = 30
    e.freeze_panes = "A2"
    e.append([])
    e.append(["Schedule B6: patient encounters in which a unit gave naloxone or recorded an opioid poisoning code (T40.0 to "
              "T40.4 or T40.6 with a sixth character of 1 to 4; adverse effects and underdosing are not poisonings), the patient "
              "a county resident and alive when EMS care ended, by the quarter of the 911 call. An encounter is one patient at "
              "one CAD incident: a unit's report that hands the patient to another unit joins the receiving unit's report (the "
              "fire districts' reports give the age by year of birth, so it can read a year above the ambulance's)."])
    e.cell(e.max_row, 1).font = Font(italic=True, size=9, color="5D6670")
    n = wb.create_sheet("Notes")
    for k, v in (("Prepared by", "%s, %s" % (W.PEOPLE["manager"], W.TITLES["manager"])),
                 ("Date", "2 June 2027"),
                 ("Death file", "State Office of Vital Records death records as extracted on 11 May 2027, with the interstate exchange receipts"),
                 ("EMS reports", "Wendmoor County EMS Data System, reports for 1 January 2021 to 31 December 2026"),
                 ("Certification", "The outcome and the phase two decision are in %s" % CERT)):
        n.append([k, v])
    n.column_dimensions["A"].width = 14
    n.column_dimensions["B"].width = 110
    for row in n.iter_rows():
        row[0].font = Font(bold=True)
    wb.save(path)


def main(target, out, figures_path):
    os.makedirs(out, exist_ok=True)
    f, ex = figures(target)
    write_certification(f, ex, os.path.join(out, CERT))
    import postprocess as PP
    PP.normalize_pdf(os.path.join(out, CERT), b"D:20270601161205-07'00'", b"D:20270602083344-07'00'")
    write_workbook(f, os.path.join(out, WORKBOOK))
    with open(figures_path, "w") as fh:
        json.dump(f, fh, indent=1, sort_keys=True)
        fh.write("\n")
    return {"figures": f, "exact": ex}
