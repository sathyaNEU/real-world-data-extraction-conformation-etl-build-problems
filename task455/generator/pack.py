"""Write the pack's data files in the layouts of the systems that produce them."""
import csv
import datetime as dt

import world as W

FILES = {
    "deaths": "wendmoor_death_records_2019_2026.csv",
    "receipts": "interstate_exchange_receipts_2019_2027.csv",
    "status": "vital_records_data_year_status.csv",
    "review": "ome_drug_death_review_2025.xlsx",
    "agreement": "ori_grant_agreement_2024.pdf",
    "ems": "ems_patient_care_reports_2021_2026.csv",
    "dict": "extract_data_dictionary.xlsx",
    "manifest": "extract_manifest_2027-05-11.csv",
    "pop": "wendmoor_population_estimates_2019_2026.csv",
    "naloxone": "ori_naloxone_kit_distribution_2024_2026.csv",
}
DEATH_COLS = ["state_file_no", "record_version", "record_status", "registered_date", "date_of_death", "sex", "age_years",
              "marital_status", "education", "residence_county", "residence_county_code", "residence_state",
              "occurrence_county", "occurrence_county_code", "occurrence_state", "resident_status", "place_of_death",
              "manner_of_death", "certifier", "underlying_cause", "record_axis"]
RECEIPT_COLS = ["state_file_no", "record_version", "sending_state", "sending_state_file_no", "transaction", "received_date"]
REVIEW_COLS = ["review_id", "state_file_no", "case_number", "date_of_death", "age", "sex", "cause_of_death",
               "drugs_caused_or_contributed", "drugs_present_not_contributing", "opioid_caused_or_contributed", "manner",
               "review_completed", "reviewer"]


def _w(path, header, rows, crlf=True):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=header, lineterminator="\r\n" if crlf else "\n")
        w.writeheader()
        for r in rows:
            w.writerow(r)
    return len(rows)


def write_deaths(rows, path):
    return _w(path, DEATH_COLS, rows)


def write_receipts(rows, path):
    rows = sorted(rows, key=lambda r: (r["received_date"], r["state_file_no"], r["record_version"]))
    return _w(path, RECEIPT_COLS, rows, crlf=False)


def write_status(path):
    notes = {2020: "Close delayed by the pandemic backlog of pending investigations",
             2022: "Close delayed by the migration to the new death registration system (live 1 January 2023)",
             2026: "Final statistical file in preparation"}
    rows = [{"data_year": y, "file_status": "Closed", "closed_date": W.CLOSE[y].isoformat(),
             "final_statistical_file_released": W.RELEASED[y].isoformat() if W.RELEASED[y] else "",
             "note": notes.get(y, "")} for y in W.YEARS]
    return _w(path, ["data_year", "file_status", "closed_date", "final_statistical_file_released", "note"], rows)


ABOUT = [
    ("Review", "State Office of the Medical Examiner review of drug overdose deaths, Wendmoor County residents, 2025"),
    ("Requested by", "Wendmoor County Department of Public Health, Epidemiology and Vital Statistics, under the Overdose Reduction Initiative grant agreement"),
    ("Deaths reviewed", "Every death on the department's list of 6 May 2026: deaths of Wendmoor County residents in 2025 whose underlying cause on the State Office of Vital Records' 2025 file, as closed on 30 April 2026, was a drug poisoning (ICD-10 X40-X44, X60-X64, X85, Y10-Y14), wherever the death occurred."),
    ("Method", "For each death a forensic pathologist of this office read the investigating office's case file, the autopsy or external examination report and the full toxicology report with drug concentrations, and recorded the drugs that caused or contributed to the death and the drugs present that did not. The review did not re-certify any death and did not change the state's file."),
    ("Records", "One row per death, keyed by the state file number on the department's list. Drug names are as the toxicology report gives them; metabolites are reported under the parent drug."),
    ("Toxicology", "State Public Health Laboratory, Forensic Toxicology Section, under the Wendmoor County Coroner's contract since March 2018 (comprehensive postmortem panel, unchanged since)."),
    ("Signed", "Leah Jackson, MD, Associate Medical Examiner, 12 February 2027"),
]


def write_review(rows, path):
    import openpyxl
    from openpyxl.styles import Alignment, Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cases"
    ws.append(REVIEW_COLS)
    for r in rows:
        ws.append([int(r[c]) if c == "age" else (r[c] or None) for c in REVIEW_COLS])
    for c in ws[1]:
        c.font = Font(bold=True)
    widths = {"A": 15, "B": 14, "C": 14, "D": 13, "E": 6, "F": 5, "G": 62, "H": 44, "I": 30, "J": 12, "K": 22, "L": 15, "M": 9}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    ws.freeze_panes = "A2"
    ab = wb.create_sheet("About")
    for k, v in ABOUT:
        ab.append([k, v])
    ab.column_dimensions["A"].width = 18
    ab.column_dimensions["B"].width = 110
    for row in ab.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    return len(rows)


EMS_COLS = ["pcr_number", "agency_name", "unit", "incident_number", "psap_call_datetime", "unit_arrived_datetime",
            "patient_age", "patient_age_units", "patient_gender", "patient_home_county", "patient_home_state",
            "primary_impression", "secondary_impression", "naloxone_administered", "naloxone_doses", "naloxone_response",
            "disposition", "destination"]


def write_ems(rows, path):
    return _w(path, EMS_COLS, rows)


AGES = ["0-4", "5-9", "10-14", "15-19", "20-24", "25-29", "30-34", "35-39", "40-44", "45-49", "50-54", "55-59", "60-64",
        "65-69", "70-74", "75-79", "80-84", "85+"]
AGE_SHARE = [.059, .062, .064, .063, .064, .070, .071, .068, .063, .059, .060, .063, .062, .055, .044, .031, .020, .022]


def write_population(path):
    """The state demographer's July 1 estimates by age group and sex (the latest series, every year)."""
    import numpy as np
    g = np.random.default_rng([W.SEED, 31])
    rows = []
    for y in W.YEARS:
        shift = (y - 2019) * .0012
        share = np.array(AGE_SHARE) * np.array([1 - shift * (i < 9) + shift * 1.4 * (i >= 13) for i in range(18)])
        share = share / share.sum()
        tot = W.POP[y]
        for i, a in enumerate(AGES):
            n = tot * share[i]
            male = .505 if i < 13 else (.47 if i < 16 else .40)
            m = int(round(n * male + g.normal(0, 25)))
            f = int(round(n * (1 - male) + g.normal(0, 25)))
            rows.append({"year": y, "estimate_date": "%d-07-01" % y, "age_group": a, "sex": "Male", "population": m})
            rows.append({"year": y, "estimate_date": "%d-07-01" % y, "age_group": a, "sex": "Female", "population": f})
    return _w(path, ["year", "estimate_date", "age_group", "sex", "population"], rows)


def write_naloxone(path):
    import numpy as np
    g = np.random.default_rng([W.SEED, 41])
    chans = [("Pharmacy standing order", 310), ("Community distribution site", 520), ("County jail at release", 85),
             ("EMS leave-behind", 140), ("Mail order", 95), ("Treatment clinic", 120)]
    rows = []
    for y in (2024, 2025, 2026):
        for m in range(1, 13):
            ramp = min(1.0, .55 + .05 * ((y - 2024) * 12 + m))
            for ch, base in chans:
                kits = int(round(base * ramp * (1 + .12 * g.standard_normal())))
                rows.append({"month": "%d-%02d" % (y, m), "channel": ch, "kits_distributed": max(0, kits),
                             "doses_per_kit": 2})
    return _w(path, ["month", "channel", "kits_distributed", "doses_per_kit"], rows)


DICT = {
    FILES["deaths"]: [
        ("state_file_no", "State file number of the death record (out-of-state deaths of residents: the number assigned when the record was received)."),
        ("record_version", "Version of the record: 0 the original registration, then each amendment or void in order."),
        ("record_status", "Original, Amendment or Void. A void withdraws the record (a duplicate registration)."),
        ("registered_date", "Date the version was registered."),
        ("date_of_death", "Date of death; the data year is the year of death."),
        ("sex", "M or F."),
        ("age_years", "Age at death in completed years."),
        ("marital_status", "Marital status at death."),
        ("education", "Highest education completed."),
        ("residence_county", "County of residence (decedent's usual residence). Blank for a residence outside the state."),
        ("residence_county_code", "Three-digit county code of the county of residence."),
        ("residence_state", "Postal code of the state of residence when outside this state; blank for this state."),
        ("occurrence_county", "County where the death occurred. Blank for a death outside the state."),
        ("occurrence_county_code", "Three-digit county code of the county of occurrence."),
        ("occurrence_state", "Postal code of the state where the death occurred when outside this state; blank for this state."),
        ("resident_status", "NCHS resident status: 1 resident (died in the county of residence), 2 intrastate nonresident, 3 interstate nonresident, 4 foreign resident."),
        ("place_of_death", "Place of death as certified."),
        ("manner_of_death", "Manner of death as certified (Natural, Accident, Suicide, Homicide, Pending investigation, Could not be determined)."),
        ("certifier", "Type of certifier: Physician, Coroner or Medical examiner."),
        ("underlying_cause", "ICD-10 underlying cause of death, as coded by the State Office of Vital Records (R99 while a cause is pending investigation)."),
        ("record_axis", "ICD-10 multiple causes of death (record axis), up to 20 codes separated by spaces. Versions registered before 1 January 2023, in the legacy system, store codes without the decimal point."),
    ],
    FILES["receipts"]: [
        ("state_file_no", "State file number of the record (as in the death file)."),
        ("record_version", "Version of the record received (as in the death file)."),
        ("sending_state", "Postal code of the state that sent the record."),
        ("sending_state_file_no", "The sending state's certificate number for the death."),
        ("transaction", "New record or Amendment, as the sending state sent it."),
        ("received_date", "Date this office received the record or amendment from the sending state."),
    ],
    FILES["status"]: [
        ("data_year", "Data year (year of death)."),
        ("file_status", "Status of the data year's file."),
        ("closed_date", "Date the State Office of Vital Records closed the data year's file."),
        ("final_statistical_file_released", "Date the final statistical file for the data year was released."),
        ("note", "Office note."),
    ],
    FILES["review"]: [
        ("review_id", "Review case identifier."),
        ("state_file_no", "State file number of the death record reviewed."),
        ("case_number", "Investigating office's case number."),
        ("date_of_death", "Date of death."),
        ("age", "Age at death in years."),
        ("sex", "M or F."),
        ("cause_of_death", "Cause of death as the reviewing pathologist states it."),
        ("drugs_caused_or_contributed", "Drugs the review found caused or contributed to the death, separated by semicolons."),
        ("drugs_present_not_contributing", "Drugs the toxicology found that the review found did not contribute, separated by semicolons."),
        ("opioid_caused_or_contributed", "Y when an opioid is among the drugs that caused or contributed, else N."),
        ("manner", "Manner of death."),
        ("review_completed", "Date the review of the case was completed."),
        ("reviewer", "Reviewing pathologist (staff code)."),
    ],
    FILES["ems"]: [
        ("pcr_number", "Patient care report number (agency code, year, sequence)."),
        ("agency_name", "EMS agency that completed the report."),
        ("unit", "Responding unit."),
        ("incident_number", "Incident number assigned by the Wendmoor Communications Center."),
        ("psap_call_datetime", "Date and time the 911 call was received (local time)."),
        ("unit_arrived_datetime", "Date and time the unit arrived on scene (local time)."),
        ("patient_age", "Patient's age."),
        ("patient_age_units", "Units of patient_age."),
        ("patient_gender", "Male or Female."),
        ("patient_home_county", "Patient's home county as recorded; blank when not obtained."),
        ("patient_home_state", "Patient's home state when outside this state; blank otherwise."),
        ("primary_impression", "Provider's primary impression (ICD-10-CM)."),
        ("secondary_impression", "Provider's secondary impression (ICD-10-CM)."),
        ("naloxone_administered", "Yes when the unit gave naloxone."),
        ("naloxone_doses", "Doses of naloxone the unit gave."),
        ("naloxone_response", "Patient response to naloxone as recorded."),
        ("disposition", "Incident/patient disposition (NEMSIS eDisposition.12)."),
        ("destination", "Receiving facility when transported; the receiving unit when care was transferred to another EMS unit."),
    ],
    FILES["pop"]: [
        ("year", "Estimate year."), ("estimate_date", "Reference date of the estimate."), ("age_group", "Five-year age group."),
        ("sex", "Male or Female."), ("population", "Estimated resident population of Wendmoor County."),
    ],
    FILES["naloxone"]: [
        ("month", "Month of distribution."), ("channel", "Distribution channel (Schedule A1)."),
        ("kits_distributed", "Naloxone kits distributed."), ("doses_per_kit", "Naloxone doses in a kit."),
    ],
}
CODE_GROUPS = [
    ("Drug poisoning (overdose) underlying causes", "X40-X44 unintentional; X60-X64 suicide; X85 homicide; Y10-Y14 undetermined intent"),
    ("Opioids", "T40.0 opium; T40.1 heroin; T40.2 other opioids (natural and semisynthetic); T40.3 methadone; T40.4 synthetic opioids other than methadone (fentanyl, tramadol); T40.6 other and unspecified narcotics"),
    ("Cocaine", "T40.5"),
    ("Psychostimulants with abuse potential", "T43.6 (methamphetamine, amphetamine)"),
    ("Cannabis", "T40.7"),
    ("Benzodiazepines", "T42.4"),
    ("Other antiepileptic and sedative-hypnotic drugs", "T42.6 (gabapentin, zolpidem)"),
    ("Antidepressants", "T43.0 tricyclic; T43.2 other antidepressants"),
    ("Other antipsychotics", "T43.5"),
    ("Antihistamines", "T45.0"),
    ("Nonopioid analgesics", "T39.1 (acetaminophen)"),
    ("Other and unspecified drugs", "T50.9"),
    ("Alcohol", "T51.0 (an additional code; not a drug poisoning on its own)"),
]


def write_dictionary(path):
    import openpyxl
    from openpyxl.styles import Alignment, Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Files"
    ws.append(["file", "sheet", "contents"])
    sheet = {FILES["deaths"]: "Death records", FILES["receipts"]: "Exchange receipts", FILES["status"]: "Data year status",
             FILES["review"]: "ME review 2025", FILES["ems"]: "EMS reports", FILES["pop"]: "Population",
             FILES["naloxone"]: "Naloxone kits"}
    desc = {
        FILES["deaths"]: "State Office of Vital Records death records: deaths of Wendmoor County residents wherever they occurred and deaths occurring in Wendmoor County, data years 2019 to 2026, every registered version, as extracted on 11 May 2027.",
        FILES["receipts"]: "Records and amendments of Wendmoor County residents' deaths in other states received through the interstate exchange, 2019 to 11 May 2027.",
        FILES["status"]: "Close dates of each data year's death file.",
        FILES["review"]: "State Office of the Medical Examiner review of the 2025 drug overdose deaths of Wendmoor County residents (Schedule B3).",
        FILES["agreement"]: "Overdose Reduction Initiative grant agreement, with Schedules A and B.",
        FILES["ems"]: "Wendmoor County EMS Data System patient care reports with naloxone given, a primary or secondary impression coded in ICD-10-CM T36 to T65, or an overdose, intoxication, altered mental status or unresponsiveness impression, 1 January 2021 to 31 December 2026.",
        FILES["pop"]: "July 1 population estimates for Wendmoor County by age group and sex, 2019 to 2026, from the state demographer's current series.",
        FILES["naloxone"]: "Overdose Reduction Initiative naloxone kits distributed by month and channel, 2024 to 2026.",
        FILES["dict"]: "This dictionary.",
        FILES["manifest"]: "Files in the 11 May 2027 extract with their sources and record counts.",
    }
    for f in FILES.values():
        ws.append([f, sheet.get(f), desc[f]])
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 110
    for name, fields in DICT.items():
        sh = wb.create_sheet(sheet[name])
        sh.append(["field", "description"])
        for k, v in fields:
            sh.append([k, v])
        sh.column_dimensions["A"].width = 32
        sh.column_dimensions["B"].width = 120
    cg = wb.create_sheet("Code groups")
    cg.append(["group", "codes"])
    for k, v in CODE_GROUPS:
        cg.append([k, v])
    cg.column_dimensions["A"].width = 46
    cg.column_dimensions["B"].width = 120
    for sh in wb.worksheets:
        for c in sh[1]:
            c.font = Font(bold=True)
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        sh.freeze_panes = "A2"
    wb.save(path)
    return sum(len(v) for v in DICT.values())


def write_manifest(path, n):
    src = {
        "deaths": ("State Office of Vital Records, death registration system (legacy system for versions before 2023)", "2027-05-11"),
        "receipts": ("State Office of Vital Records, interstate exchange log", "2027-05-11"),
        "status": ("State Office of Vital Records, data year status", "2027-05-11"),
        "review": ("State Office of the Medical Examiner", "2027-02-12"),
        "agreement": ("Wendmoor Valley Recovery Foundation and Wendmoor County", "2024-01-22"),
        "ems": ("Wendmoor County EMS Data System", "2027-05-07"),
        "pop": ("State demographer, county population estimates by age and sex", "2027-04-02"),
        "naloxone": ("Department of Public Health, Overdose Reduction Initiative program records", "2027-01-15"),
        "dict": ("Department of Public Health, Epidemiology and Vital Statistics", "2027-05-11"),
    }
    rows = []
    for k, f in FILES.items():
        if k == "manifest":
            continue
        s, d = src[k]
        rows.append({"file": f, "source": s, "extracted": d, "records": n.get(k, "")})
    return _w(path, ["file", "source", "extracted", "records"], rows)
