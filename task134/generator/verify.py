"""Independent verifier: reads only the shipped bundle (target/) and recomputes every graded figure with its own code,
then compares with figures.json. Imports nothing from the generator.

    python3 task134/generator/verify.py [--target DIR]
"""
import argparse
import datetime as dt
import json
import math
import os
import sys

import numpy as np
import openpyxl
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
EPOCH = dt.date(2000, 1, 1)
M = 30.4375
ASOF = (dt.date(2026, 12, 31) - EPOCH).days
LO = (dt.date(2021, 4, 1) - EPOCH).days
QE = [(d - EPOCH).days for d in (dt.date(2027, 6, 30), dt.date(2027, 9, 30), dt.date(2027, 12, 31), dt.date(2028, 3, 31))]
LOCS = ["Allanvale", "Braeburn", "Carsefield", "Dunmorrow", "Kirkhallam"]
# the charging policy (3.2): the upper capital limit rose from 28,750 to 33,750 on 6 April 2026 and rises to 40,750 on
# 5 April 2027 (checked against the document's text below)
S1, S2 = (dt.date(2026, 4, 6) - EPOCH).days, (dt.date(2027, 4, 5) - EPOCH).days
RISE1, RISE2 = 33750 - 28750, 40750 - 33750


def d(s):
    return float((dt.date.fromisoformat(s) - EPOCH).days) if s else math.inf


def read(tdir):
    rd = lambda f, sep=",": pd.read_csv(os.path.join(tdir, f), sep=sep, dtype=str, keep_default_na=False)
    X = dict(pl=rd("care_home_placements_2021-2026.csv"), li=rd("fpnc_payment_lines_2021-2026.csv"),
             cw=rd("client_id_crosswalk_2023.csv"), dl=rd("nrs_death_linkage_2027-01-15.tsv", "\t"),
             fa=rd("charging_assessments_2021-2026.csv"), lk=rd("postcode_locality_lookup.csv"))
    ws = openpyxl.load_workbook(os.path.join(tdir, "care_home_register_2026-12.xlsx"), read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    X["reg"] = pd.DataFrame([dict(zip(rows[0], r)) for r in rows[1:]])
    return X


def people(X):
    who = {}
    refs = {}
    for r in X["cw"].itertuples(index=False):
        p = r.merged_into or r.client_id
        who[r.client_id] = p
        if r.legacy_client_no:
            who[r.legacy_client_no] = p
        if r.debtor_ref:
            refs[r.debtor_ref] = p
    died = {r.client_id: d(r.date_of_death) for r in X["dl"].itertuples(index=False)}
    return who, refs, died


def stays(X, who, died):
    pl = X["pl"][X["pl"]["placement_type"] == "PERM"].copy()
    pl["p"] = pl["client_id"].map(lambda c: who.get(c, c))
    pl["s"] = pl["start_date"].map(d)
    pl["e"] = pl["end_date"].map(d)
    out = []
    for p, g in pl.sort_values(["p", "s"]).groupby("p"):
        cur = None
        for r in g.itertuples(index=False):
            if cur and r.s == cur["e"] + 1 and cur["why"] in ("FEE", "MOVE", "SVC"):
                cur.update(e=r.e, why=r.end_reason, last=r)
                continue
            if cur:
                out.append(cur)
            cur = dict(p=p, adm=d(r.admitted), s=r.s, e=r.e, why=r.end_reason, route=r.route, eps=[], last=r)
        out.append(cur)
    eps = {}
    for r in pl.itertuples(index=False):
        eps.setdefault(r.p, []).append(r)
    for c in out:
        if c["why"] in ("DIED", "DISCH"):
            c["x"] = c["e"]
        else:
            c["x"] = died[c["p"]] if c["p"] in died and c["s"] <= died[c["p"]] <= ASOF else math.inf
        c["eps"] = [r for r in eps[c["p"]] if c["s"] <= r.s and (r.e <= c["e"] or math.isinf(c["e"]))]
    return out


def awards(X, refs, died):
    out = []
    for ref, g in X["li"].groupby("debtor_ref"):
        r = g.iloc[0]
        if r["end_reason"] == "CF":
            end, why = d(r["award_end"]) + 1, "C"
        elif r["end_reason"] in ("DIED", "LEFT"):
            end, why = d(r["award_end"]), "X"
        else:
            p = refs.get(ref)
            end, why = (died[p], "X") if p in died and died[p] <= ASOF else (math.inf, "")
        out.append((d(r["admitted"]), end, why))
    return out


def rates(rows, nb=180):
    """rows: (admission day, at-risk from, at-risk to, event?) -> occurrence / exposure per month by month since
    admission, in the window 1 April 2021 to 31 December 2026"""
    expo = np.zeros(nb)
    occ = np.zeros(nb)
    for a, s, e, ev in rows:
        s2, e2 = max(s, LO), min(e, ASOF)
        if e2 <= s2:
            continue
        m0, m1 = (s2 - a) / M, (e2 - a) / M
        b = min(int(m0), nb - 1)
        while b < nb:
            hi = b + 1 if b < nb - 1 else 1e9
            lo = max(m0, b)
            if lo >= m1:
                break
            expo[b] += min(m1, hi) - lo
            b += 1
        if ev and LO < e <= ASOF:
            occ[min(int(m1), nb - 1)] += 1
    return np.divide(occ, expo, out=np.zeros(nb), where=expo > 0)


def cumh(h):
    c = np.concatenate([[0.0], np.cumsum(h)])

    def H(x):
        x = np.asarray(x, float)
        b = np.minimum(x.astype(int), len(h) - 1)
        return c[b] + h[b] * (x - b)
    return H


def forecast(X):
    who, refs, died = people(X)
    st = stays(X, who, died)
    aw = awards(X, refs, died)
    hx = rates([(c["adm"], c["s"], c["x"], not math.isinf(c["x"])) for c in st])
    hc = rates([(a, a, e, w == "C") for a, e, w in aw])
    hs = rates([(a, a, e, w == "X") for a, e, w in aw])
    H = cumh(hx)
    q = np.array([(x - ASOF) / M for x in QE])
    now = [c for c in st if c["s"] <= ASOF < c["x"]]
    d0 = np.array([(ASOF - c["adm"]) / M for c in now])
    surv = np.array([np.exp(-(H(d0 + qq) - H(d0))).sum() for qq in q])
    lo3, hi3 = (dt.date(2024, 1, 1) - EPOCH).days, (dt.date(2027, 1, 1) - EPOCH).days
    rate = sum(1 for c in st if c["route"] == "DIRECT" and lo3 <= c["s"] < hi3) / 36.0
    direct = np.zeros(4)
    for k in range(4):
        n = int(math.ceil(q[k] / 0.1 - 1e-9))
        w = q[k] / n
        tin = (np.arange(n) + 0.5) * w
        direct[k] = rate * w * np.exp(-(H(q[k] - tin) - H(0.0))).sum()
    # the 6 April 2026 rise: that day's conversions against an ordinary week's (the 52 weeks before)
    cd = np.array([e for a, e, w in aw if w == "C"])
    on_day = int((cd == S1).sum())
    flow = ((cd >= S1 - 364) & (cd < S1)).sum() / 52.0
    w1 = (on_day - flow / 7.0) / flow
    w2 = w1 * RISE2 / RISE1
    members = [(a, e) for a, e, w in aw if e > ASOF and a <= ASOF]
    pool = np.array([(ASOF - a) / M for a, e in members])
    sh = np.array([w1 * 7 / M if a < S1 else 0.0 for a, e in members])
    HC = cumh(hc)
    tA = (S2 - ASOF) / M
    a2 = w2 * 7 / M
    conv = np.zeros(4)
    s = np.ones(len(pool))
    t = 0.0
    jumped = False
    while t < q[3] - 1e-9:
        if not jumped and t >= tA - 1e-9:
            cdur = pool + sh + t
            pj = s * (1 - np.exp(-(HC(cdur + a2) - HC(cdur))))
            for k in range(4):
                if t < q[k]:
                    conv[k] += (pj * np.exp(-(H(pool + q[k]) - H(pool + t)))).sum()
            s = s - pj
            sh = sh + a2
            jumped = True
        hh = min(0.1, q[3] - t)
        if not jumped and t + hh > tA:
            hh = tA - t
        for qq in q:
            if t < qq - 1e-9 < t + hh:
                hh = qq - t
        b = np.minimum((pool + sh + t).astype(int), 179)
        bx = np.minimum((pool + t).astype(int), 179)
        lc, lx = hc[b], hs[bx]
        tot = lc + lx
        p = s * np.divide(lc, tot, out=np.zeros_like(tot), where=tot > 0) * (1 - np.exp(-tot * hh))
        mid = t + hh / 2
        for k in range(4):
            if mid < q[k]:
                conv[k] += (p * np.exp(-(H(pool + q[k]) - H(pool + mid)))).sum()
        s = s * np.exp(-tot * hh)
        t += hh
    since = {"1 January %d" % y: sum(1 for a, e, w in aw if e > ASOF and (dt.date(y, 1, 1) - EPOCH).days <= a <= ASOF)
             for y in (2025, 2026)}
    parts = dict(today_residents=surv, direct_admissions=direct, conversions=conv, rate=rate, since=since,
                 on_day=on_day, flow=flow, w1=w1, w2=w2)
    return surv + direct + conv, len(now), len(pool), st, parts


def rates_2026(tdir):
    """the 2026-27 rows of the contract rate table: the interim rate the case system holds and the agreed rate"""
    ws = openpyxl.load_workbook(os.path.join(tdir, "nchc_rates_2021-22_to_2026-27.xlsx"), read_only=True).active
    interim, agreed = {}, {}
    for r in ws.iter_rows(values_only=True):
        if r and r[0] == "2026-27":
            (interim if str(r[4]).startswith("Interim") else agreed)[r[1]] = float(r[2])
    return interim, agreed


def ask(X, st, tdir):
    who, refs, died = people(X)
    interim, agreed = rates_2026(tdir)
    lk = dict(zip(X["lk"]["postcode"], X["lk"]["locality"]))
    home = {r["service_no"]: lk.get(r["postcode"]) for r in X["reg"].to_dict("records")}
    byp = {}
    for r in X["cw"].itertuples(index=False):
        if r.debtor_ref:
            byp.setdefault(r.merged_into or r.client_id, []).append(r.debtor_ref)
    fa = X["fa"][X["fa"]["outcome"] != "SELF"]
    fa = fa.assign(e=fa["effective_from"].map(d), c=fa["completed"].map(d))
    fa = fa[fa["e"] <= ASOF]
    best = {}
    for r in fa.sort_values(["e", "c"]).itertuples(index=False):
        best[r.debtor_ref] = (r.e, r.c, float(r.weekly_contribution) if r.weekly_contribution else 0.0)
    cells = {}
    for c in st:
        if not (c["s"] <= ASOF < c["x"]):
            continue
        ep = [r for r in c["eps"] if r.s <= ASOF and (math.isinf(r.e) or r.e >= ASOF)][-1]
        loc = home.get(ep.service_no)
        if loc is None:
            continue
        assert ep.start_date >= "2026-04-01"
        fee = float(ep.weekly_fee) - interim[ep.care_type] + agreed[ep.care_type]
        cand = [best[x] for x in byp.get(c["p"], []) if x in best]
        con = max(cand)[2] if cand else 0.0
        cells.setdefault((loc, ep.care_type), []).append((fee, fee - con))
    return {k: (np.mean([a for a, b in v]), np.mean([b for a, b in v]), len(v)) for k, v in cells.items()}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--target", default=os.path.join(TASK, "target"))
    a = ap.parse_args()
    X = read(a.target)
    F = json.load(open(os.path.join(HERE, "figures.json")))
    fails = 0
    tot, ncf, npool, st, parts = forecast(X)
    for k, qlab in enumerate(["30 June 2027", "30 September 2027", "31 December 2027", "31 March 2028"]):
        want = F["call"][qlab]
        good = abs(tot[k] - want["unrounded"]) < 0.01 and int(math.floor(tot[k] / 10 + 0.5) * 10) == want["stated"]
        fails += not good
        print("%s  %-18s %.2f (figures.json %.2f, stated %d)" % ("ok  " if good else "FAIL", qlab, tot[k], want["unrounded"], want["stated"]))
        for part in ("today_residents", "direct_admissions", "conversions"):
            good = abs(parts[part][k] - F["components"][qlab][part]) < 0.06
            fails += not good
            print("%s    %-18s %.2f (figures.json %.1f)" % ("ok  " if good else "FAIL", part, parts[part][k], F["components"][qlab][part]))
    good = ncf == F["stock_31_december_2026"] and npool == F["pool_31_december_2026"]
    fails += not good
    print("%s  stock %d, pool %d" % ("ok  " if good else "FAIL", ncf, npool))
    good = abs(parts["rate"] - F["direct_admissions_per_month"]) < 0.006 and parts["since"] == F["pool_admitted_since"]
    fails += not good
    print("%s  direct admissions %.3f a month; pool admitted since %s" % ("ok  " if good else "FAIL", parts["rate"], parts["since"]))
    good = parts["on_day"] == F["stage_1"]["conversions_that_day"] and abs(parts["flow"] - F["stage_1"]["ordinary_week"]) < 1e-3 \
        and abs(parts["w1"] - F["stage_1"]["weeks"]) < 0.006 and abs(parts["w2"] - F["stage_2"]["weeks"]) < 0.006
    fails += not good
    print("%s  6 April 2026: %d conversions, an ordinary week %.3f, %.2f weeks; 5 April 2027 %.2f weeks"
          % ("ok  " if good else "FAIL", parts["on_day"], parts["flow"], parts["w1"], parts["w2"]))
    from pypdf import PdfReader
    pol = " ".join(p.extract_text() for p in PdfReader(os.path.join(a.target, "charging_policy_2026-27.pdf")).pages)
    pol = " ".join(pol.split())
    good = all(x in pol for x in ("£33,750", "£28,750", "£40,750", "6 April 2026", "5 April 2027"))
    fails += not good
    print("%s  the charging policy states the limits the forecast uses" % ("ok  " if good else "FAIL"))
    A = ask(X, st, a.target)
    for loc in LOCS:
        for ct in ("Residential", "Nursing"):
            f, n, cnt = A[(loc, ct)]
            want = F["ask"]["%s %s" % (loc, ct.lower())]
            good = abs(f - want["fee"]) < 1e-3 and abs(n - want["net"]) < 1e-3 and cnt == want["residents"] \
                and format(int(math.floor(f + 0.5)), ",") == want["fee_stated"] and format(int(math.floor(n + 0.5)), ",") == want["net_stated"]
            fails += not good
            print("%s  %-10s %-11s fee %.2f net %.2f residents %d" % ("ok  " if good else "FAIL", loc, ct, f, n, cnt))
    print("verify: %d failed" % fails)
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
