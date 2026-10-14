"""The documents and workbooks: the order, the target methodology, the branch's reference tables, Schedule A, the
zoning schedule, the city's extract notes, and two files the work does not use (population projections and the
rental market survey)."""
import csv
import datetime as dt
import os
import shutil
import tempfile

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

import officex as OX
import postprocess as PP
import world as W

P = W.PEOPLE


def _pdf(path, blocks, header, footer, props, seed, doc_id, stamp, margins=(1418, 1418, 1247, 1247)):
    tmp = tempfile.mkdtemp()
    src = os.path.join(tmp, "doc.docx")
    OX.docx(src, blocks, header, footer, props, font="Calibri", seed=seed, sessions=6, margins=margins)
    OX.pdf_from_docx(src, path)
    shutil.rmtree(tmp)
    PP.normalize_pdf(path, stamp, doc_id=doc_id)


def order_pdf(path):
    p = lambda t: ("p", t)
    blocks = [
        ("title", "Transit-Oriented Area Order (Ridgemont Line), 2026"),
        ("small", "Province of British Columbia  |  %s  |  Order made 30 June 2026" % W.MINISTRY),
        p("The Minister of Housing, under the transit-oriented area provisions of the Local Government Act, designates the "
          "land around the four Ridgemont Line stations in the City of Calderbrook as transit-oriented areas and orders "
          "as follows."),
        ("h1", "1. Interpretation"),
        p("In this order, \"lot\" means a parcel listed in Schedule A, and includes each parcel it becomes after the date "
          "of this order; \"floor area ratio\" means the gross floor area of the buildings on a lot divided by the area of "
          "the lot; \"exchange point\" means the point at each station from which distances in this order are measured."),
        ("h1", "2. Transit-oriented areas"),
        p("The transit-oriented areas are the land within 800 metres of the exchange points of Ellery, Westerly, Haldane and "
          "Pell Road stations. Each lot is in the distance band of its nearest point to the exchange point: within 200 "
          "metres, 200 to 400 metres, or 400 to 800 metres."),
        ("h1", "3. Minimum allowable density"),
        p("On a lot where the City's zoning permits residential use, the City must permit residential development to at "
          "least the floor area ratio and height set out for the lot's distance band:"),
        ("table", [["Distance band", "Floor area ratio, at least", "Height, at least"],
                   ["Within 200 metres", "5.0", "20 storeys"],
                   ["200 to 400 metres", "4.0", "12 storeys"],
                   ["400 to 800 metres", "3.0", "8 storeys"]], [3600, 2800, 2600], 19),
        p("Where the City's zoning permits a greater floor area ratio or height on a lot, the City's zoning applies to that "
          "extent."),
        ("h1", "4. Off-street parking"),
        p("The City must not require off-street parking spaces for residential use on a lot, other than spaces for persons "
          "with disabilities."),
        ("h1", "5. Schedule A"),
        p("Schedule A lists the lots by folio as at the date of this order, with each lot's station and distance band. "
          "Schedule A is published with this order as a workbook."),
        ("h1", "6. Effective date"),
        p("This order takes effect on 1 January 2027."),
    ]
    props = dict(title="Transit-Oriented Area Order (Ridgemont Line), 2026", creator=P["branch_lead"],
                 modified_by=P["branch_lead"], created=dt.datetime(2026, 5, 11, 10, 4),
                 modified=dt.datetime(2026, 6, 30, 15, 22), revision=11, total_time=212, company="Province of British Columbia")
    _pdf(path, blocks, ("Province of British Columbia", "Transit-Oriented Area Order (Ridgemont Line)"),
         ("Ministry of Housing  |  30 June 2026", "Page {PAGE} of {NUMPAGES}"), props, 2606,
         b"3B7E19C04A2D5F86E1C09A4B72D3E618", dt.datetime(2026, 6, 30, 15, 31))


def methodology_pdf(path):
    p = lambda t: ("p", t)
    blocks = [
        ("title", "Housing Target Methodology 2026"),
        ("small", "%s  |  %s  |  March 2026" % (W.MINISTRY, W.BRANCH)),
        ("h1", "1. Purpose"),
        p("This methodology sets out what a housing target counts. It applies to the five-year targets the Minister sets "
          "for municipalities under the Housing Supply Act, including a target for the lots in a transit-oriented area "
          "order. It replaces the 2023 methodology for targets set after 1 April 2026."),
        ("h1", "2. Targets for a transit-oriented area order"),
        p("2.1 A target counts the lots listed in the order's Schedule A, each carried to the parcel or parcels it becomes "
          "after the Schedule's date."),
        p("2.2 A target is the number of net new homes that building permits issued in each calendar year of the target "
          "period will add on those lots: the dwelling units in building permits issued for new buildings, less the "
          "dwelling units standing on the lots that the permitted work clears. A dwelling unit is a self-contained home; "
          "a house with a secondary suite or a laneway house holds two. Dwelling units cleared are counted once for each "
          "permit for new buildings, whether that permit, a separate demolition permit or both record them."),
        p("2.3 A permit is counted in the calendar year it is issued. A permit that is cancelled or expires is not counted. "
          "A permit that revises an earlier permit for the same buildings replaces it: the buildings are counted once, in "
          "the year the earlier permit was issued, with the dwelling units of the latest revision."),
        p("2.4 Secondary suites and laneway houses added to an existing home are counted under section 4 and are not part "
          "of a transit-oriented area target."),
        p("2.5 A target is set on the zoning in force when the order takes effect: the order's minimum densities on the "
          "lots in Schedule A, or the municipality's zoning where it permits more. No other rezoning is assumed during the "
          "target period."),
        p("2.6 The target period is the five calendar years from the order's effective date. A target states the total "
          "for the period and the figure for each year."),
        ("h1", "3. Municipal records"),
        p("Municipalities provide their building permit, occupancy and folio records to 31 December each year by the "
          "following 15 January, with the assessment roll extract the branch requests."),
        ("h1", "4. Secondary suites and laneway houses"),
        p("Secondary suites and laneway houses are counted toward a municipality's citywide target from the year their "
          "permits are issued, at one dwelling unit each, and are reported separately from any transit-oriented area "
          "target."),
        ("h1", "5. Monitoring"),
        p("5.1 Municipalities report permits issued and homes completed on the lots in a transit-oriented area order each "
          "year."),
        p("5.2 A home is completed in the year the building it is in is first granted occupancy, interim or final."),
        p("5.3 The homes standing on a lot are the dwelling units in the buildings on it at the assessment roll's condition "
          "date, 31 October of the year before the roll year."),
    ]
    props = dict(title="Housing Target Methodology 2026", creator=P["branch_lead"], modified_by=P["analyst"],
                 created=dt.datetime(2026, 1, 19, 9, 37), modified=dt.datetime(2026, 3, 24, 16, 8), revision=17,
                 total_time=361, company="Province of British Columbia")
    _pdf(path, blocks, (W.MINISTRY, "Housing Target Methodology 2026"),
         (W.BRANCH + "  |  March 2026", "Page {PAGE} of {NUMPAGES}"), props, 2603,
         b"9D04E2B7716A3C58F0E14D29B6A8C372", dt.datetime(2026, 3, 24, 16, 15))


def _header(ws, row, n, fill="DDEBF7"):
    for j in range(1, n + 1):
        c = ws.cell(row=row, column=j)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(vertical="top", wrap_text=True)


def reference_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Building types"
    ws.append(["Building types used in housing target work"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    hdr = ["Type", "Description", "Construction", "Storeys", "Floor area ratio on a typical lot", "Net to gross",
           "Average home (net sq ft)"]
    ws.append(hdr)
    _header(ws, 3, len(hdr))
    rows = [("TH3", "Townhouses and rowhouses", "Wood frame", "3", "up to 1.2", 0.92, 1350),
            ("LR4", "Low-rise apartment", "Wood frame", "4", "up to 2.0", 0.85, 820),
            ("MR6", "Mid-rise apartment", "Wood frame", "5 to 6", "3.0", 0.85, 820),
            ("HR", "High-rise apartment", "Concrete", "7 and over", "above 3.0, to the permitted maximum", 0.82, 790)]
    for r in rows:
        ws.append(list(r))
    ws.append([])
    ws.append(["Gross floor area is the lot area times the floor area ratio. Net floor area is gross floor area times "
               "net to gross. Homes in a building: its net floor area divided by the average home, rounded down to a "
               "whole home."])
    for col, w in zip("ABCDEFG", [8, 26, 14, 11, 34, 12, 14]):
        ws.column_dimensions[col].width = w
    for r in range(4, 8):
        ws.cell(row=r, column=6).number_format = "0.00"
        ws.cell(row=r, column=7).number_format = "#,##0"

    ws = wb.create_sheet("Sale values")
    ws.append(["New home sale values, $ per net square foot"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    hdr = ["Year", "Area", "Townhouse", "Wood frame apartment", "Concrete apartment"]
    ws.append(hdr)
    _header(ws, 3, len(hdr))
    for y in W.YEARS:
        for a in ["ELL", "WES", "HAL", "PEL", "CC", "NC", "RAV", "FER"]:
            ws.append([y, W.AREA_NAME[a], W.value(a, "town", y), W.value(a, "wood", y), W.value(a, "conc", y)])
    for r in range(4, ws.max_row + 1):
        for j in (3, 4, 5):
            ws.cell(row=r, column=j).number_format = "#,##0"
    for col, w in zip("ABCDE", [8, 20, 12, 21, 19]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"

    ws = wb.create_sheet("Construction costs")
    ws.append(["All-in development cost, $ per gross square foot"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    hdr = ["Year", "TH3", "LR4", "MR6", "HR"]
    ws.append(hdr)
    _header(ws, 3, len(hdr))
    for y in W.YEARS:
        ws.append([y] + [W.cost(t, y) for t in ("TH3", "LR4", "MR6", "HR")])
    for r in range(4, ws.max_row + 1):
        for j in range(2, 6):
            ws.cell(row=r, column=j).number_format = "#,##0"
    for col, w in zip("ABCDE", [8, 10, 10, 10, 10]):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet("Notes")
    notes = [
        "Housing Targets Branch planning reference tables. Updated 20 November 2026 with the 2026 market review and cost guide.",
        "",
        "Sale values: average first-sale prices of new homes per net square foot, by area and product, from Halvorsen Land "
        "Economics' annual new home market review. Each year's figures cover first sales in that calendar year (2026: sales "
        "to 30 September 2026). Areas: the four Ridgemont Line station areas (within 800 metres of each station), the City "
        "Centre, and the neighbourhoods of North Calder, Ravensdale and Fernlea.",
        "",
        "Construction costs: all-in development cost per gross square foot by building type, from Brockmore Cost Consulting's "
        "annual construction cost guide for the Calderbrook market: hard and soft costs, financing, development cost charges "
        "and permit fees, and the developer's margin. Each year's figures are the guide's edition for that year.",
        "",
        "Building types: the branch's standard types for the region. Floor area ratios are those each type reaches on a "
        "typical lot of the size found in the region's single-detached, corridor and apartment zones.",
        "",
        "Contact: %s, Housing Targets Branch." % P["analyst"],
    ]
    for t in notes:
        ws.append([t])
    ws.column_dimensions["A"].width = 120
    for r in range(1, ws.max_row + 1):
        ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    OX.xlsx_from_openpyxl(wb, path, P["analyst"], P["analyst"], dt.datetime(2023, 2, 6, 9, 48),
                          dt.datetime(2026, 11, 20, 14, 37), seed=4101)


def schedule_xlsx(rows, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Schedule A"
    ws.append(["Transit-Oriented Area Order (Ridgemont Line), 2026: Schedule A, lots as at 30 June 2026"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    hdr = ["Folio", "Civic address", "Station", "Distance band", "Floor area ratio, at least", "Height, at least (storeys)"]
    ws.append(hdr)
    _header(ws, 3, len(hdr))
    for r in rows:
        addr = ("%s-%d %s" % (r["unit"], r["civic"], r["street"])) if r["unit"] else "%d %s" % (r["civic"], r["street"])
        ws.append([r["folio"], addr, W.AREA_NAME[r["station"]], W.TIER_BAND[r["tier"]], W.TIER_FAR[r["tier"]],
                   W.TIER_STOREYS[r["tier"]]])
    for r in range(4, ws.max_row + 1):
        ws.cell(row=r, column=5).number_format = "0.0"
    for col, w in zip("ABCDEF", [17, 30, 12, 16, 14, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = "A3:F%d" % ws.max_row
    OX.xlsx_from_openpyxl(wb, path, P["branch_lead"], P["branch_lead"], dt.datetime(2026, 5, 4, 11, 12),
                          dt.datetime(2026, 6, 30, 14, 58), seed=4102)


ZONING_ROWS = [
    # code, name, residential use, max FAR, from, to, bylaw, note
    ("RS-1", "Single detached residential", "Single detached house, with a secondary suite or a laneway house", "0.70",
     "2001-03-12", "2020-03-01", "1450", "Renamed R1 in the Zoning Bylaw consolidation"),
    ("RT-2", "Townhouse infill residential", "Townhouses and rowhouses", "1.20", "2019-06-10", "2020-03-01", "19-212",
     "Townhouse infill, North Calder, Ravensdale and Fernlea; renamed RM1 in the consolidation"),
    ("C-2", "Community commercial", "Not permitted", "", "2001-03-12", "2020-03-01", "1450",
     "Commercial floor area only; renamed C2 in the consolidation"),
    ("RM-3", "Low-rise apartment residential", "Apartments to 4 storeys", "1.50", "2001-03-12", "2020-03-01", "1450",
     "Renamed RA1 in the consolidation"),
    ("RM-5", "Apartment and townhouse residential", "Apartments to 4 storeys and townhouses", "1.80", "2001-03-12",
     "2020-03-01", "1450", "Renamed RA3 in the consolidation"),
    ("CR-3", "Corridor mixed use", "Apartments to 6 storeys, commercial at grade permitted", "3.00", "2013-04-15",
     "2020-03-01", "13-044", "Arterial corridors; renamed MU3 in the consolidation"),
    ("C-3", "Town centre commercial", "Apartments to 6 storeys above commercial", "3.00", "2001-03-12", "2016-04-17",
     "1450", "Replaced by CD-CC on the City Centre lots"),
    ("CD-CC", "City Centre comprehensive development", "Apartments, no height limit", "5.50", "2016-04-18", "2020-03-01",
     "16-077", "City Centre plan; renamed CC1 in the consolidation"),
    ("R1", "Single detached residential", "Single detached house, with a secondary suite or a laneway house", "0.70",
     "2020-03-02", "", "2020-001", ""),
    ("RM1", "Townhouse residential", "Townhouses and rowhouses", "1.20", "2020-03-02", "", "2020-001", ""),
    ("C2", "Community commercial", "Not permitted", "", "2020-03-02", "", "2020-001", "Commercial floor area only"),
    ("RA1", "Low-rise apartment", "Apartments to 4 storeys", "1.50", "2020-03-02", "", "2020-001", ""),
    ("RA3", "Apartment and townhouse", "Apartments to 4 storeys and townhouses", "1.80", "2020-03-02", "", "2020-001", ""),
    ("MU3", "Corridor mixed use", "Apartments to 6 storeys, commercial at grade permitted", "3.00", "2020-03-02", "",
     "2020-001", ""),
    ("CC1", "City Centre", "Apartments, no height limit", "5.50", "2020-03-02", "", "2020-001", ""),
]


def zoning_csv(path):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(["zone_code", "zone_name", "residential_use", "max_floor_area_ratio", "in_force_from", "in_force_to",
                    "bylaw", "note"])
        for r in ZONING_ROWS:
            w.writerow(r)


def notes_txt(path, counts):
    t = """CITY OF CALDERBROOK
Planning and Development Services: data extracts for the Ministry of Housing
Ridgemont Line transit-oriented area target, 2027 to 2031
Prepared 12 January 2027 by %(gis)s, GIS and Data Services. Permit records: %(permits)s, Building Division.

FILES
  roll_extract_2014-2018.csv          assessment roll extract, roll years 2014 to 2018 (former layout), %(r1)s records
  roll_extract_2019-2026.csv          assessment roll extract, roll years 2019 to 2026 (current layout), %(r2)s records
  building_permits_2014-2026.csv      building permits issued 1 January 2014 to 31 December 2026, %(bp)s records
  permit_buildings.csv                buildings on the permits for new construction, %(pb)s records
  occupancy_permits_2018-2026.csv     occupancy permits issued 1 January 2018 to 31 December 2026, %(oc)s records
  folio_history_2014-2026.csv         folios retired and created, 2014 to 2026, %(fh)s records
  zoning_schedule_2014-2026.csv       zone codes in force since 2014, with residential use and maximum floor area ratio
The permit, occupancy and folio files were pulled from the permit system on 11 January 2027.

SCOPE
Every folio within 800 metres of Ellery, Westerly, Haldane and Pell Road stations, every folio in the City Centre,
and the folios in the corridor and townhouse infill zones of North Calder, Ravensdale and Fernlea.
Within 800 metres of the stations each strata lot is its own record. Elsewhere a strata plan is carried as one
record for the plan (use STRATA PLAN SUMMARY / Strata plan summary) with its number of strata lots.

ROLL EXTRACTS
One record per folio per roll year. A roll year's values are as at 1 July of the year before, and its property
records as at 31 October of the year before. Values are in dollars.
The 2019 roll introduced the current layout. Former columns and their current equivalents:
  FOLIO = folio; CIVIC_ADDRESS = civic_number, street_name and unit_no; NEIGHBOURHOOD = neighbourhood;
  ZONE = zone_code; ACTUAL_USE = use_group; LAND_AREA, AREA_UNIT = lot_area, area_unit;
  LAND_VALUE = assessed_land; IMPROVEMENT_VALUE = assessed_improvements; YEAR_BUILT = year_built;
  STRATA_PLAN = strata_plan; FINISHED_AREA = floor_area_sqft.
The current layout adds dwelling_units and strata_lot_count; the former layout records suites and laneway
houses in ACTUAL_USE.
Lot area is in square feet, or in acres where the area unit is AC. Strata lots carry no lot area.
zone_code is the zone at the valuation date, as the zoning bylaw then named it (codes changed in the March 2020
consolidation; see zoning_schedule_2014-2026.csv).
dwelling_units: the dwelling units in a house or a purpose-built rental building. floor_area_sqft: the finished
floor area of a strata lot where the assessor has recorded it.

BUILDING PERMITS
One record per permit. folio is the folio the permit was issued against. permit_type: New Construction,
Demolition, Secondary Suite, Laneway House, Alteration/Repair.
floor_area is gross floor area: square feet on permits issued to 31 December 2019 and square metres from
1 January 2020, when the permit system moved to metric; floor_area_unit says which.
dwelling_units_created and dwelling_units_demolished are as entered on each permit.
status at the pull: Issued, Completed (final occupancy granted), Cancelled, Expired, Superseded. A permit issued
to revise another names it in revises_permit.

PERMIT BUILDINGS
One record per building on a permit for new construction: storeys, construction, dwelling units and gross floor
area (in the permit's floor area unit).

OCCUPANCY PERMITS
One record per occupancy permit, interim or final, against a building permit and building number.
dwelling_units are as entered on the occupancy permit.

FOLIO HISTORY
One record for each retired folio and each folio created from it, with the event and the plan deposited.
""" % dict(gis=P["city_gis"], permits=P["permits"], **counts)
    with open(path, "w", encoding="utf-8", newline="\r\n") as f:
        f.write(t)


def projections_csv(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["area", "year", "population", "households", "persons_per_household"])
        for r in rows:
            w.writerow(r)


def rental_xlsx(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "October survey"
    ws.append(["Rental Market Survey, October: purpose-built rental apartments, Calderbrook survey zones"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Universe: privately initiated purpose-built rental structures of three or more units. Rents are average "
               "asking rents for occupied and vacant units, $ a month."])
    ws.append([])
    hdr = ["Survey year", "Zone", "Bedroom type", "Units in universe", "Vacancy rate (%)", "Average rent ($)"]
    ws.append(hdr)
    _header(ws, 4, len(hdr), fill="E2EFDA")
    for r in rows:
        ws.append(list(r))
    for r in range(5, ws.max_row + 1):
        ws.cell(row=r, column=4).number_format = "#,##0"
        ws.cell(row=r, column=5).number_format = "0.0"
        ws.cell(row=r, column=6).number_format = "#,##0"
    for col, w in zip("ABCDEF", [12, 28, 14, 16, 16, 16]):
        ws.column_dimensions[col].width = w
    OX.xlsx_from_openpyxl(wb, path, "Housing Market Analysis", "Housing Market Analysis", dt.datetime(2026, 12, 3, 8, 15),
                          dt.datetime(2026, 12, 14, 10, 2), seed=4103, company="")


def index_txt(path):
    rows = [
        ("tod_area_order_ridgemont_2026.pdf", "Ministry of Housing, the order as made", "2026-06-30", "the order, sections 1 to 6"),
        ("tod_order_schedule_a.xlsx", "Ministry of Housing, Schedule A as made", "2026-06-30", "lots in the four transit-oriented areas"),
        ("housing_target_methodology_2026.pdf", "Housing Targets Branch", "2026-03-24", "what a housing target counts"),
        ("housing_reference_tables_2026.xlsx", "Housing Targets Branch", "2026-11-20", "building types; sale values and costs, 2014 to 2026"),
        ("roll_extract_2014-2018.csv", "City of Calderbrook, GIS and Data Services", "2027-01-08", "assessment roll, roll years 2014 to 2018"),
        ("roll_extract_2019-2026.csv", "City of Calderbrook, GIS and Data Services", "2027-01-08", "assessment roll, roll years 2019 to 2026"),
        ("building_permits_2014-2026.csv", "City of Calderbrook, Building Division", "2027-01-11", "building permits issued 2014 to 2026"),
        ("permit_buildings.csv", "City of Calderbrook, Building Division", "2027-01-11", "buildings on permits for new construction"),
        ("occupancy_permits_2018-2026.csv", "City of Calderbrook, Building Division", "2027-01-11", "occupancy permits issued 2018 to 2026"),
        ("folio_history_2014-2026.csv", "City of Calderbrook, Building Division", "2027-01-11", "folios retired and created, 2014 to 2026"),
        ("zoning_schedule_2014-2026.csv", "City of Calderbrook, Planning", "2026-12-15", "zone codes in force since 2014"),
        ("extract_notes_2027-01-12.txt", "City of Calderbrook, GIS and Data Services", "2027-01-12", "notes on the city's extracts"),
        ("population_projections_2026-2036.csv", "Provincial statistics office, 2026 release", "2026-09-17", "population and households by area, 2026 to 2036"),
        ("rental_market_survey_oct2026.xlsx", "Rental Market Survey, October 2026 release", "2026-12-14", "purpose-built rental vacancy and rents, 2022 to 2026"),
    ]
    w = max(len(r[0]) for r in rows) + 2
    w2 = max(len(r[1]) for r in rows) + 2
    lines = ["HOUSING TARGETS BRANCH, MINISTRY OF HOUSING",
             "City of Calderbrook: Ridgemont Line transit-oriented area target, 2027 to 2031. Working folder index.",
             "Kept by %s. Last updated 15 January 2027." % P["analyst"], "",
             "File".ljust(w) + "From".ljust(w2) + "Dated       Covers"]
    for f, src, d, cov in rows:
        lines.append(f.ljust(w) + src.ljust(w2) + d + "  " + cov)
    lines += ["",
              "The City's files were provided under section 3 of the Housing Target Methodology for target work. The order, "
              "Schedule A, the methodology and the reference tables are released under the Open Government Licence - British "
              "Columbia; the projections and the rental survey under their publishers' open licences."]
    with open(path, "w", encoding="utf-8", newline="\r\n") as f:
        f.write("\n".join(lines) + "\n")
