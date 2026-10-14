"""The answer from the shipped pack: figures.json for the write-up, and the two golden deliverables.

    python3 task134/generator/golden.py
"""
import datetime as dt
import json
import os
import sys
import warnings

import numpy as np

warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import estimators as E
import officex as OX
import world as W

TASK = os.path.dirname(HERE)
TARGET = os.path.join(TASK, "target")
GOLDEN = os.path.join(TASK, "golden")
KIND_LABEL = {"house": "Detached houses", "comm": "Commercial sites", "rental": "Purpose-built rental buildings"}


def compute():
    pack = E.load(TARGET)
    P = E.panel(pack)
    F = E.forward(pack)
    G = E.golden(pack, P, F)
    yr, tot = E.rnd(G["yearly"])
    r = G["rates"]
    net = (F["homes"] - F["units"])
    by_kind = {}
    for k in ("house", "comm", "rental"):
        m = F["kind"] == k
        by_kind[k] = dict(lots=int(m.sum()), rate_pct=round(100 * r[k], 1), rate=r[k], events=r["e"][k], lot_years=r["n"][k],
                          net_homes=int(net[m].sum()))
    a = E.ask_a(pack)
    b = E.ask_b(pack)
    d = E.default_path(pack)
    riv = E.rivals(pack, P, F, G)
    fam = E.family(pack, P, F, G)
    fig = {
        "call": {"total": tot, "yearly": {str(y): v for y, v in zip(W.FWD, yr)}},
        "unrounded": {"total": round(G["total"], 1), "yearly": {str(y): round(float(v), 1) for y, v in zip(W.FWD, G["yearly"])}},
        "components": {
            "standing_lots": int(len(F)),
            "standing_by_kind": {k: v["lots"] for k, v in by_kind.items()},
            "rates_pct": {k: v["rate_pct"] for k, v in by_kind.items()},
            "rates_exact": {k: round(v["rate"], 5) for k, v in by_kind.items()},
            "rebuilds": {k: v["events"] for k, v in by_kind.items()},
            "feasible_lot_years": {k: v["lot_years"] for k, v in by_kind.items()},
            "net_homes_once_rebuilt": int(net.sum()),
            "net_homes_by_kind": {k: v["net_homes"] for k, v in by_kind.items()},
            "best_building": "MR6 at 3.0 FAR on every standing lot",
            "standing_lots_r_min": round(float(F["R"].min()), 2),
            "lots_2026_permits_excluded": int(lots_permitted_2026(pack)),
        },
        "order": {"band_m": [200, 400, 800], "far": [5.0, 4.0, 3.0], "storeys": [20, 12, 8],
                  "mr6_far": 3.0, "net_to_gross": 0.85, "home_sqft": 820},
        "ask_a": {s: {str(y): a[(s, y)] for y in E.ASK_YEARS} for s in E.STATIONS},
        "ask_a_totals": {"by_year": {str(y): sum(a[(s, y)] for s in E.STATIONS) for y in E.ASK_YEARS},
                         "by_station": {s: sum(a[(s, y)] for y in E.ASK_YEARS) for s in E.STATIONS},
                         "all": sum(a.values())},
        "ask_b": {s: {t: b[(s, t)] for t in E.BANDS} for s in E.STATIONS},
        "ask_b_totals": {"by_station": {s: sum(b[(s, t)] for t in E.BANDS) for s in E.STATIONS},
                         "by_band": {t: sum(b[(s, t)] for s in E.STATIONS) for t in E.BANDS},
                         "all": sum(b.values())},
        "default": {"takeup_pct": round(100 * d["takeup"], 2), "capacity": int(round(d["capacity"])),
                    "yearly": int(round(d["yearly"][0])), "yearly_100": int(round(d["yearly"][0] / 100.0) * 100),
                    "total": int(round(d["total"])), "total_500": int(round(d["total"] / 500.0) * 500),
                    "corridor_capacity": int(round(d["corridor_capacity"])), "corridor_homes": d["corridor_homes"]},
        "rivals": {k: {"yearly": [int(round(x)) for x in v], "total": int(round(float(v.sum()))),
                       "total_500": int(round(float(v.sum()) / 500.0) * 500)} for k, v in riv.items()},
        "family": {k: {"yearly": [round(float(x), 1) for x in v], "total": round(float(v.sum()), 1)} for k, v in fam.items()},
    }
    return pack, P, F, G, fig


def lots_permitted_2026(pack):
    ev = E.events(pack)
    r26 = pack["roll"][pack["roll"]["year"] == 2026].set_index("folio")
    cur = E.current_folios(pack)
    n = 0
    for f in cur["folio"]:
        if ev.get(f) == 2026 and f in r26.index and isinstance(r26.loc[f, "kind"], str):
            n += 1
    return n


def write_xlsx(fig, path):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = openpyxl.Workbook()
    thin = Side(style="thin", color="A6A6A6")
    head_fill = PatternFill("solid", fgColor="DCE6F1")

    def header(ws, row, n):
        for j in range(1, n + 1):
            c = ws.cell(row=row, column=j)
            c.font = Font(bold=True)
            c.fill = head_fill
            c.border = Border(bottom=thin)
            c.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.active
    ws.title = "Target 2027-31"
    ws.append(["%s net new homes on the Ridgemont Line order's lots over 2027 to 2031, falling from %s in 2027 to %s in 2031"
               % (format(fig["call"]["total"], ","), format(fig["call"]["yearly"]["2027"], ","),
                  format(fig["call"]["yearly"]["2031"], ","))])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Net new homes that building permits will add on the lots in Schedule A of the Transit-Oriented Area Order "
               "(Ridgemont Line), carried to their current folios (Housing Target Methodology 2026, section 2)."])
    ws.append([])
    ws.append(["Year", "Net new homes"])
    header(ws, 4, 2)
    for y in W.FWD:
        ws.append([y, fig["call"]["yearly"][str(y)]])
    ws.append(["2027 to 2031", fig["call"]["total"]])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
    for r in range(5, ws.max_row + 1):
        ws.cell(row=r, column=2).number_format = "#,##0"
    ws.append(["Yearly figures to the nearest 100 homes, the total to the nearest 500, each rounded from the unrounded "
               "estimate."])
    ws.cell(row=ws.max_row, column=1).font = Font(italic=True, size=9, color="595959")
    ws.append([])
    ws.append(["Basis"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    c = fig["components"]
    lines = [
        "Standing lots: the houses, commercial sites and purpose-built rental buildings on the order's lots at the 2026 roll "
        "with no counted permit for new construction to 31 December 2026: %s lots." % format(c["standing_lots"], ","),
        "Building that pays best under the order: the six-storey wood frame at 3.0 FAR on every standing lot. Concrete "
        "high-rise costs more than it sells for at all four station areas, so the 4.0 and 5.0 FAR bands add nothing over 3.0.",
        "Feasibility: that building's residual land value (lot area x FAR x (net to gross x sale value - all-in cost), from "
        "the 2026 reference tables) exceeds the assessed land and improvements on every standing lot.",
        "Rebuild rates: permits for new construction on feasible lots of each kind, 2014 to 2026, over the feasible "
        "lot-years: houses %.1f%%, commercial sites %.1f%%, rental buildings %.1f%% a year." % (
            c["rates_pct"]["house"], c["rates_pct"]["comm"], c["rates_pct"]["rental"]),
        "Each lot is rebuilt once, so each year applies the rate to the lots still standing; homes are net of the dwelling "
        "units on the 2026 roll (%s net homes once every standing lot is rebuilt)." % format(c["net_homes_once_rebuilt"], ","),
        "The %d lots whose permits for new construction were issued in 2026 count in 2026's permits and are not among "
        "the standing lots." % c["lots_2026_permits_excluded"],
    ]
    for t in lines:
        ws.append([t])
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=6)
        ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[ws.max_row].height = 31
    ws.append([])
    ws.append(["Kind of lot", "Standing lots", "Rebuild rate (% a year)", "Net homes once rebuilt"])
    header(ws, ws.max_row, 4)
    for k in ("house", "comm", "rental"):
        ws.append([KIND_LABEL[k], c["standing_by_kind"][k], c["rates_pct"][k], c["net_homes_by_kind"][k]])
        ws.cell(row=ws.max_row, column=3).number_format = "0.0"
        ws.cell(row=ws.max_row, column=2).number_format = "#,##0"
        ws.cell(row=ws.max_row, column=4).number_format = "#,##0"
    for col, w in zip("ABCDEF", [30, 16, 16, 18, 14, 14]):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet("Completions 2022-26")
    by_year = fig["ask_a_totals"]["by_year"]
    peak = max(by_year, key=lambda y: by_year[y])
    ws.append(["Completions on the order's lots peaked at %s homes in %s and were %s in 2026" % (
        format(by_year[peak], ","), peak, format(by_year["2026"], ","))])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Homes completed on the order's lots by station area and year: each building's dwelling units in the year it "
               "was first granted occupancy, interim or final (methodology 5.2)."])
    ws.append([])
    ws.append(["Station area"] + [str(y) for y in E.ASK_YEARS] + ["2022 to 2026"])
    header(ws, 4, 7)
    ws.freeze_panes = "B5"
    for s in E.STATIONS:
        ws.append([s] + [fig["ask_a"][s][str(y)] for y in E.ASK_YEARS] + [fig["ask_a_totals"]["by_station"][s]])
    ws.append(["All four"] + [fig["ask_a_totals"]["by_year"][str(y)] for y in E.ASK_YEARS] + [fig["ask_a_totals"]["all"]])
    for j in range(1, 8):
        ws.cell(row=ws.max_row, column=j).font = Font(bold=True)
    for r in range(5, ws.max_row + 1):
        for j in range(2, 8):
            ws.cell(row=r, column=j).number_format = "#,##0"
    for col, w in zip("ABCDEFG", [16, 10, 10, 10, 10, 10, 14]):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet("Homes standing 2026")
    ws.append(["%s homes stood on the order's lots at the 2026 roll" % format(fig["ask_b_totals"]["all"], ",")])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Homes standing by station area and distance band: the dwelling units in the buildings on each lot at the "
               "roll's condition date, 31 October 2025 (methodology 5.3). Parking and storage stalls registered as strata "
               "lots are not homes."])
    ws.append([])
    ws.append(["Station area"] + E.BANDS + ["All bands"])
    header(ws, 4, 5)
    ws.freeze_panes = "B5"
    for s in E.STATIONS:
        ws.append([s] + [fig["ask_b"][s][t] for t in E.BANDS] + [fig["ask_b_totals"]["by_station"][s]])
    ws.append(["All four"] + [fig["ask_b_totals"]["by_band"][t] for t in E.BANDS] + [fig["ask_b_totals"]["all"]])
    for j in range(1, 6):
        ws.cell(row=ws.max_row, column=j).font = Font(bold=True)
    for r in range(5, ws.max_row + 1):
        for j in range(2, 6):
            ws.cell(row=r, column=j).number_format = "#,##0"
    for col, w in zip("ABCDE", [16, 15, 16, 16, 12]):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet("Notes")
    notes = [
        ("As of", "Records to 31 December 2026: the 2026 assessment roll (values at 1 July 2025, property records at "
                  "31 October 2025); building, occupancy and folio records pulled from the City of Calderbrook permit "
                  "system on 11 January 2027; the branch's reference tables updated 20 November 2026."),
        ("Lots", "Schedule A of the Transit-Oriented Area Order (Ridgemont Line), 2026, as at 30 June 2026, each folio "
                 "carried to the folios it became through the City's folio history."),
        ("Net new homes", "Dwelling units in permits issued for new buildings, less the dwelling units standing on the "
                          "lots they clear (methodology 2.2); a permit counts in the year it is issued, a revision "
                          "once (2.3). Secondary suites and laneway houses are outside the target (2.4)."),
        ("Zoning", "The order's minimum densities from 1 January 2027, or the City's zoning where it permits more; no "
                   "other rezoning (2.5)."),
        ("Completed", "A home is completed in the year its building is first granted occupancy, interim or final (5.2)."),
        ("Rounding", "Target years to the nearest 100 homes and the total to the nearest 500, each from the unrounded "
                     "estimate; every other count a whole number."),
        ("Prepared by", "%s, %s, 15 January 2027." % (W.PEOPLE["analyst"], W.BRANCH)),
    ]
    ws.append(["Notes"])
    ws["A1"].font = Font(bold=True, size=12)
    for k, v in notes:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=1).alignment = Alignment(vertical="top")
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[ws.max_row].height = 46
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 100
    OX.xlsx_from_openpyxl(wb, path, W.PEOPLE["analyst"], W.PEOPLE["analyst"], dt.datetime(2027, 1, 13, 9, 12),
                          dt.datetime(2027, 1, 15, 16, 41), seed=4201)


def write_png(fig, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "svg.hashsalt": "task134",
                         "path.simplify": True})
    colours = {"Ellery": "#1F4E79", "Westerly": "#2E75B6", "Haldane": "#9DC3E6", "Pell Road": "#C9A227"}
    years = E.ASK_YEARS
    fig_, ax = plt.subplots(figsize=(9.0, 5.6), dpi=110)
    bottom = np.zeros(len(years))
    for s in E.STATIONS:
        vals = np.array([fig["ask_a"][s][str(y)] for y in years], float)
        bars = ax.bar([str(y) for y in years], vals, bottom=bottom, color=colours[s], edgecolor="white", linewidth=0.8,
                      label=s, width=0.62)
        for k, (b_, v) in enumerate(zip(bottom, vals)):
            if v >= 45:
                ax.text(k, b_ + v / 2, format(int(v), ","), ha="center", va="center", fontsize=8.5,
                        color="white" if s in ("Ellery", "Westerly") else "#1A1A1A")
            elif v > 0:
                ax.annotate(format(int(v), ","), xy=(k + 0.31, b_ + v / 2), xytext=(k + 0.42, b_ + v / 2 + 30),
                            fontsize=8, va="center", ha="left", color="#1A1A1A",
                            arrowprops=dict(arrowstyle="-", color="#7F7F7F", lw=0.6))
        bottom += vals
    for k, t in enumerate(bottom):
        ax.text(k, t + 25, format(int(t), ","), ha="center", va="bottom", fontsize=9, fontweight="bold")
    ax.set_ylabel("Homes completed (first occupancy)")
    by_year = fig["ask_a_totals"]["by_year"]
    peak = max(by_year, key=lambda y: by_year[y])
    ax.set_title("Completions on the Ridgemont Line order's lots peaked at %s homes in %s" % (format(by_year[peak], ","), peak),
                 fontsize=11.5, loc="left")
    ax.spines[["top", "right"]].set_visible(False)
    ax.set_ylim(0, max(bottom) * 1.12)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, p: format(int(v), ",")))
    h_, l_ = ax.get_legend_handles_labels()
    ax.legend(h_[::-1], l_[::-1], title="Station area", frameon=False, loc="upper left", bbox_to_anchor=(1.0, 1.0))
    fig_.text(0.01, 0.01, "Source: City of Calderbrook occupancy and building permits, Schedule A lots carried through the "
              "folio history.", fontsize=7.5, color="#595959")
    fig_.tight_layout(rect=(0, 0.03, 1, 1))
    fig_.savefig(path, metadata={"Software": None})
    plt.close(fig_)


def main():
    os.makedirs(GOLDEN, exist_ok=True)
    pack, P, F, G, fig = compute()
    with open(os.path.join(HERE, "figures.json"), "w", encoding="utf-8") as f:
        json.dump(fig, f, indent=1)
        f.write("\n")
    write_xlsx(fig, os.path.join(GOLDEN, "tod_housing_target_2027-31.xlsx"))
    write_png(fig, os.path.join(GOLDEN, "tod_completions_2022-26.png"))
    print("call", fig["call"], "\nunrounded", fig["unrounded"])


if __name__ == "__main__":
    main()
