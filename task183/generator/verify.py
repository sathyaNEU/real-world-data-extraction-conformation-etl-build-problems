"""Independent verifier: reads only the shipped files in target/ (and the goldens and figures.json it checks), with
its own code path, no generator module and no generator state. Recomputes the call (each in-service section's
rate for the cable failures it has already had, estimated on section days since 2012), the default path's
feeder, and every graded ask figure: SAIFI, SAIDI and the four-or-more count for the eight feeders."""
import csv
import datetime as dt
import json
import math
import os
import sys
from collections import defaultdict
from zoneinfo import ZoneInfo

import openpyxl
import pyarrow.parquet as pq

EASTERN = ZoneInfo("America/New_York")
UTC = dt.timezone.utc
CUT = dt.date(2026, 9, 30)
START = dt.date(2012, 1, 1)
RY = (dt.date(2025, 10, 1), dt.date(2026, 9, 30))


def day(s):
    if "/" in s:
        m, d, y = s.split("/")
        return dt.date(int(y), int(m), int(d))
    return dt.date(int(s[:4]), int(s[5:7]), int(s[8:10]))


def call(target):
    regwb = openpyxl.load_workbook(os.path.join(target, "gis_ug_primary_sections_20260930.xlsx"), read_only=True)
    rows = list(regwb.active.iter_rows(values_only=True))
    h = [i for i, r in enumerate(rows) if r and r[0] == "SECTION_ID"][0]
    cols = [c.lower() for c in rows[h]]
    secs = {}
    for r in rows[h + 1:]:
        if not r or not r[0]:
            continue
        d = dict(zip(cols, r))
        to_date = lambda x: None if x is None else (x.date() if isinstance(x, dt.datetime) else x)
        secs[d["section_id"]] = (d["feeder_id"], d["status"], to_date(d["install_date"]), to_date(d["retired_date"]))
    fails = defaultdict(list)
    recent = defaultdict(int)
    with open(os.path.join(target, "ug_cable_fault_work_orders_2012_2026.csv"), newline="") as fh:
        for r in csv.DictReader(fh):
            if r["cause_code"] not in ("INS", "SPL", "TRM"):
                continue
            d = day(r["fault_date"])
            fails[r["section_id"]].append(d)
            if dt.date(2021, 10, 1) <= d <= CUT and secs[r["section_id"]][1] == "IS":
                recent[secs[r["section_id"]][0]] += 1
    # section days at risk by failures already on the section (0, 1, 2, 3+), and the failures in each state
    days, events = defaultdict(float), defaultdict(int)
    for sid, (f, status, inst, ret) in secs.items():
        lo, hi = max(inst, START), min(ret or CUT, CUT)
        if hi <= lo:
            continue
        fl = sorted(fails.get(sid, []))
        k0 = sum(1 for d in fl if d < lo)
        t = lo
        k = k0
        for d in fl:
            if d < lo:
                continue
            days[min(k, 3)] += (d - t).days
            events[min(k, 3)] += 1
            k += 1
            t = d
        days[min(k, 3)] += (hi - t).days
    rate = {k: events[k] / (days[k] / 365.25) for k in range(4)}
    fc = defaultdict(float)
    for sid, (f, status, inst, ret) in secs.items():
        if status == "IS":
            fc[f] += rate[min(len(fails.get(sid, [])), 3)]
    return dict(fc), dict(recent), rate


def asks(target):
    devs = json.load(open(os.path.join(target, "feeder_device_list.json")))
    parent = {d["device_id"]: d["parent_device_id"] for d in devs["devices"]}
    chains = {}

    def chain(x):
        if x not in chains:
            c, y = [], x
            while y:
                c.append(y)
                y = parent[y]
            chains[x] = c
        return chains[x]
    t = pq.read_table(os.path.join(target, "connectivity_meter_points.parquet")).to_pylist()
    conn = []
    for r in t:
        ef = r["eff_from"]
        et = r["eff_to"]
        conn.append((r["meter_id"], r["device_id"], r["feeder_id"], ef, et))
    under = defaultdict(list)
    for m, dv, f, a, b in conn:
        for x in chain(dv):
            under[x].append((m, dv, f, a, b))

    def served(d):
        seen = {}
        for m, dv, f, a, b in conn:
            if a <= d and (b is None or d < b):
                seen[m] = f
        return seen

    def stamp(s, sysname):
        x = dt.datetime.strptime(s, "%Y-%m-%d %H:%M:%S")
        return x.replace(tzinfo=EASTERN) if sysname == "OMS" else x.replace(tzinfo=UTC).astimezone(EASTERN)
    outages = {}
    with open(os.path.join(target, "oms_outage_export_2020-10_2026-09.csv"), newline="") as fh:
        for r in csv.DictReader(fh):
            outages.setdefault(r["outage_id"], (r["device_id"], stamp(r["start_ts"], r["source_system"])))
    st = defaultdict(list)
    with open(os.path.join(target, "oms_restoration_steps_2020-10_2026-09.csv"), newline="") as fh:
        for r in csv.DictReader(fh):
            st[r["outage_id"]].append((int(r["step_no"]), r["restoring_device_id"],
                                       stamp(r["step_restore_ts"], r["source_system"])))
    recs = []   # (outage, meter, feeder, local day, minutes)
    for oid, (dv, start) in outages.items():
        d = start.date()
        steps = sorted(st[oid])
        for m, mdv, f, a, b in under[dv]:
            if not (a <= d and (b is None or d < b)):
                continue
            ch = chain(mdv)
            best = max((s for s in steps if s[1] in ch), key=lambda s: (len(chain(s[1])), s[2]))
            recs.append((oid, m, f, d, (best[2] - start).total_seconds() / 60))
    sustained = [x for x in recs if x[4] > 5]
    cmi = defaultdict(float)
    for _, _, _, d, mins in sustained:
        cmi[d] += mins
    logs = []
    for y in range(2021, 2026):
        n = len(served(dt.date(y, 9, 30)))
        d = dt.date(y - 1, 10, 1)
        while d <= dt.date(y, 9, 30):
            if cmi.get(d, 0) > 0:
                logs.append(math.log(cmi[d] / n))
            d += dt.timedelta(days=1)
    mu = sum(logs) / len(logs)
    sd = math.sqrt(sum((x - mu) ** 2 for x in logs) / (len(logs) - 1))
    thr = math.exp(mu + 2.5 * sd)
    now = served(RY[1])
    ntot = len(now)
    meds = {d for d, v in cmi.items() if RY[0] <= d <= RY[1] and v / ntot > thr}
    per_f = defaultdict(int)
    for f in now.values():
        per_f[f] += 1
    ci, cm = defaultdict(int), defaultdict(float)
    per_m = defaultdict(set)
    for oid, m, f, d, mins in sustained:
        if not (RY[0] <= d <= RY[1]):
            continue
        per_m[m].add(oid)
        if d in meds:
            continue
        ci[f] += 1
        cm[f] += mins
    c4 = defaultdict(int)
    for m, s in per_m.items():
        if len(s) >= 4 and m in now:
            c4[now[m]] += 1
    feeders = sorted(per_f)
    return ({f: ci[f] / per_f[f] for f in feeders}, {f: cm[f] / per_f[f] for f in feeders},
            {f: c4[f] for f in feeders}, sorted(meds), thr)


def main(target, golden, figures_path=None):
    figures_path = figures_path or os.path.join(os.path.dirname(os.path.abspath(__file__)), "figures.json")
    figs = json.load(open(figures_path))
    fc, recent, rate = call(target)
    top = max(fc, key=fc.get)
    second = sorted(fc.values())[-2]
    default = max(recent, key=recent.get)
    assert top == figs["recommended_feeder_id"] == "QH21", (top, fc)
    assert fc[top] / second >= 1.25, fc
    assert default == "ML07", recent
    saifi, saidi, c4, meds, thr = asks(target)
    n = 0
    for f in saifi:
        assert "%.2f" % saifi[f] == figs["saifi_ry2026"][f], (f, saifi[f], figs["saifi_ry2026"][f])
        assert "%.1f" % saidi[f] == figs["saidi_ry2026_min"][f], (f, saidi[f], figs["saidi_ry2026_min"][f])
        assert c4[f] == figs["customers_4plus_ry2026"][f], (f, c4[f], figs["customers_4plus_ry2026"][f])
        n += 3
    assert [d.isoformat() for d in meds] == figs["major_event_days_ry2026"], meds
    wb = openpyxl.load_workbook(os.path.join(golden, "cable_program_2027.xlsx"), read_only=True)
    assert wb["2027 program"]["B4"].value == "Quarry Hill 21"
    print("verify: call %s (%.1f against %.1f expected 2027 failures; default path %s), rates by failures so far %s; "
          "%d ask figures reproduce; major event days %s (threshold %.2f min)" % (
              top, fc[top], second, default, " ".join("%.3f" % rate[k] for k in range(4)), n,
              ", ".join(d.isoformat() for d in meds), thr))


if __name__ == "__main__":
    task = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    main(os.path.join(task, "target"), os.path.join(task, "golden"))
