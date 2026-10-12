"""Writers for the pack's record files: the work orders, the GIS section register, the code list, the outage export
and its restoration steps, the connectivity model, the device list, the customer system extract, the published
major event days, the vault inspections and the feeder loading workbook."""
import csv
import datetime as dt
import json

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import world as W

WO = "ug_cable_fault_work_orders_2012_2026.csv"
REGISTER = "gis_ug_primary_sections_20260930.xlsx"
CODES = "wo_cause_repair_codes.csv"
OMS = "oms_outage_export_2020-10_2026-09.csv"
STEPS = "oms_restoration_steps_2020-10_2026-09.csv"
CONN = "connectivity_meter_points.parquet"
DEVICES = "feeder_device_list.json"
CIS = "cis_service_accounts_20260930.csv"
MED = "major_event_days_RY2021-RY2025.csv"
VAULTS = "vault_manhole_inspections_2022-2026.csv"
LOADING = "feeder_peak_loading_2021-2026.xlsx"


def feeder_at(s, d):
    if s.lateral and d >= W.LATERAL_MOVE:
        return "ML07"
    return s.feeder


def current_feeder(s):
    return "ML07" if s.lateral else s.feeder


# ---------------------------------------------------------------- work orders
# Crew notes as typed in the field on the order (SAP PM long text, first line). Several crews, several habits: some
# write full sentences, some abbreviate everything, a few leave it blank for the supervisor to fill in. Orders keyed
# in the old work management system before the March 2014 cut-over came across in capitals, cut at 40 characters.
NOTES = {
    ("INS", "SPLICE"): [
        "insulation failure {d} ft from {a}, cut out damaged length, two cold shrink splices",
        "cable faulted to ground {d} ft from {a}; spliced in place",
        "fault {d} ft from {a}, cable wet at fault, replaced {e} ft w/ splices",
        "thumped and located {d} ft out of {a}. splice kit + {e} ft",
        "INS fault btwn {a} and {b}, repaired with splices",
        "primary fault, insulation punctured, spliced",
        "TDR {d}' from {a}, dug {e}' pit, 2 splices, re-energized",
        "loc fault w/ thumper {d}ft fr {a}. insul burnt thru. splice x2",
        "cable blew {d} ft from {b} toward {a}. new {e} ft piece, 2 cold shrinks",
        "faulted phase to gnd near {a}. treeing visible at fault. spliced, hipot ok",
        "pri cable failed {d}' out of {b}, neutral corroded at fault, spliced",
        "found fault {d} ft from {a} under driveway, bored around, 2 splices",
        "insulation breakdown approx {d} ft from {a}. splice + {e} ft cable",
        "radar/thump {a}-{b}, fault at {d} ft. cut in {e} ft w 2 splices",
        "cable failure btwn {a} & {b}, spliced",
        "",
    ],
    ("SPL", "SPLICE"): [
        "existing splice failed in {a}, new cold shrink splice",
        "tape splice failed {d} ft from {a}, re-spliced",
        "splice in {a} flashed, cut out and re-spliced",
        "failed splice {d}' from {b}, cut out, 2 new splices + {e}'",
        "old splice failed (water in splice body), new splice installed",
        "splice blew in {a}, replaced w/ heat shrink",
        "factory splice {d} ft fr {a} failed. new splice",
        "splice failure at {a}, cleaned up, new 3M QS splice",
        "spl fail in {a} - burnt, replaced",
        "found failed splice {d} ft from {b}, spliced",
        "",
    ],
    ("TRM", "TERM"): [
        "elbow failed at {a}, new 200A elbow and {e} ft",
        "termination failed at {a} riser, new term",
        "loadbreak elbow flashed over in {a}, replaced elbow",
        "term failed at {b}, re-terminated",
        "elbow burnt at {b} xfmr, new LB elbow",
        "riser term at {a} failed, new cold shrink term",
        "200A elbow failed in {b}, replaced elbow + {e} ft",
        "flashed elbow at {a}, re-termed",
    ],
    ("INS", "REPL"): [
        "insulation failure, replaced section {a} to {b} ({l} ft)",
        "faulted, new cable pulled {a}-{b}",
        "cable faulted, section replaced",
        "cable faulted {d} ft from {b}, section replaced",
        "fault {d}' from {a}. pulled new {l}' {a} to {b}, old cable abandoned",
        "new cable {a}-{b} installed after fault, energized",
        "faulted at {d} ft, replaced whole run {a} to {b}",
    ],
    ("SPL", "REPL"): [
        "splice failed, section replaced {a} to {b}",
        "splice failure, new cable pulled {a}-{b}",
        "failed splice in {a}, pulled new cable to {b} ({l}')",
    ],
    ("TRM", "REPL"): [
        "term failed at {b}, section replaced {a}-{b}",
        "elbow failed at {b}, new cable {a}-{b} + terms",
    ],
    ("TPD", "SPLICE"): [
        "dig-in by contractor (directional bore) {d} ft from {a}, spliced",
        "dig-in, fence post auger {d} ft from {a}. 2 splices",
        "third party excavation damage near {a}, spliced",
        "dig-in by landscaper, spliced, claim filed",
        "struck by water co. excavation, spliced",
        "contractor hit cable w/ backhoe {d}' from {b}, no locate ticket. 2 splices",
        "dig in - bore crew {d} ft fr {a}, spliced. claim to risk mgmt",
        "cable cut by excavator near {b}, spliced {e} ft",
        "homeowner dug into cable installing fence near {a}. spliced",
    ],
    ("OTH", "SPLICE"): [
        "no cable fault found, blown fuse traced to xfmr", "manhole fire {a}, secondary cables, primary spliced",
        "lightning damage at riser {a}, spliced", "rodent damage at {a}, spliced", "fault at {a} - pothead, repaired",
        "patrolled {a} to {b}, nothing found, fuse replaced and held", "xfmr bushing failure at {b}, cable ok",
        "squirrel in {b}, cable damaged at entry, spliced",
    ],
}
HWY_NOTES = ["SR 443 duct, cable wet, spliced in {a}", "fault under SR 443 crossing, spliced in vault {a}",
             "under highway (SR 443), spliced in place, cable full of water", "SR 443 crossing fault, re-spliced",
             "SR443 xing - water in duct, pumped {a}, spliced", "rt 443 crossing, fault {d}' from {a}, spliced",
             "443 duct bank, splice in {a}. duct full of water"]
# a field crew's description sometimes disagrees with the cause the supervisor coded (both are cable failures)
CROSS = {"INS": ["splice failed {d} ft from {a}, re-spliced", "elbow side of {b} failed, spliced"],
         "SPL": ["insulation failure near {a}, spliced", "cable faulted {d}' out of {a}, spliced"],
         "TRM": ["cable faulted at {b}, spliced and new elbow"]}
ABBREV = [(" feet", "'"), (" ft", "'"), ("approx ", "~"), (" and ", " & "), (" with ", " w/ "), ("from ", "fr "),
          ("between ", "btwn "), ("replaced", "repl"), ("installed", "inst"), ("transformer", "xfmr")]
LEGACY_UNTIL = dt.date(2014, 3, 10)


def _crew_text(rng, tmpl, a, b, d, e, l, day):
    note = tmpl.format(a=a, b=b, d=d, e=e, l=l)
    if not note:
        return note
    u = rng.random()
    if u < 0.22:
        for x, y in ABBREV:
            note = note.replace(x, y)
    elif u < 0.3:
        note = note[0].upper() + note[1:]
        if not note.endswith("."):
            note += "."
    if day < LEGACY_UNTIL:
        note = note.upper()[:40].rstrip()
    return note


def work_orders(secs, others, seed=W.SEED + 11):
    rng = np.random.default_rng(seed)
    rn = np.random.default_rng(seed + 100)       # the crew notes draw apart, so the orders and dates stay as built
    ends = structures(secs)
    events = []
    for s in secs:
        for i, (d, c) in enumerate(s.fails):
            repl = s.retired is not None and s.repl_reason == "SECOND" and i == len(s.fails) - 1
            events.append((d, s, c, repl))
    for s, d, c in others:
        events.append((d, s, c, False))
    events.sort(key=lambda e: (e[0], e[1].sid, e[2]))
    rows = []
    no = 4100082311
    for d, s, c, repl in events:
        no += int(rng.integers(140, 3900))
        if repl:
            rep = "REPL"
            done = s.retired
        elif c == "TRM":
            rep = "TERM"
            done = d + dt.timedelta(days=int(rng.integers(0, 4)))
        else:
            rep = "SPLICE"
            done = d + dt.timedelta(days=int(rng.integers(0, 9)))
        if done > W.CUT:
            done = W.CUT
        key = (c, rep) if (c, rep) in NOTES else (c, "SPLICE")
        if s.hwy and c in ("INS", "SPL") and rn.random() < 0.55:
            tmpl = HWY_NOTES[int(rn.integers(len(HWY_NOTES)))]
        elif rep != "REPL" and c in CROSS and rn.random() < 0.035:
            tmpl = CROSS[c][int(rn.integers(len(CROSS[c])))]
        else:
            tmpl = NOTES[key][int(rn.integers(len(NOTES[key])))]
        a, b = ends[s.sid]
        dd = int(rn.integers(3, max(4, s.length // 5))) * 5 if s.length > 20 else 10
        note = _crew_text(rn, tmpl, a, b, dd, int(rn.integers(2, 9)) * 5, s.length, d)
        rows.append(dict(wo_number=str(no), feeder_id=feeder_at(s, d), section_id=s.sid, fault_date=d.isoformat(),
                         completed_date=done.isoformat(), cause_code=c, repair_code=rep, crew_note=note))
    return rows


# ---------------------------------------------------------------- structures at the sections' ends
# Manholes carry the district map sheet and a number on the sheet (MH-34117 is manhole 117 on sheet 34); pad-mounted
# transformers and switchgear carry the six-digit asset tag stencilled on the cabinet. Neither has anything to do with
# the GIS facility id of the cable between them. A section replaced by new cable runs between the same structures.
SHEETS = {"HR12": (31, 32), "HR14": (31, 32), "ML03": (33, 34), "ML07": (33, 34), "QH21": (35, 36),
          "QH22": (35, 36), "SP05": (37, 38), "SP09": (37, 38)}
_ENDS = {}


def structures(secs, seed=W.SEED + 16):
    key = tuple(sorted(s.sid for s in secs))
    if key in _ENDS:
        return _ENDS[key]
    rng = np.random.default_rng(seed)
    used = set()

    def mh(f):
        while True:
            x = "MH-%d%03d" % (int(rng.choice(SHEETS[f])), int(rng.integers(1, 1000)))
            if x not in used:
                used.add(x)
                return x

    def pm():
        while True:
            x = "PM-%06d" % int(rng.integers(204000, 291000))
            if x not in used:
                used.add(x)
                return x
    by_id = {s.sid: s for s in secs}
    ends = {}
    originals = sorted([s for s in secs if s.pred is None], key=lambda s: (W.FEEDERS.index(s.feeder), s.sid))
    nodes, last_f = [], None
    for s in originals:
        if s.feeder != last_f:
            nodes, last_f = [mh(s.feeder)], s.feeder
        here = nodes[-1] if rng.random() < 0.8 else nodes[int(rng.integers(len(nodes)))]
        there = mh(s.feeder) if s.method == "CONDUIT" else pm()
        nodes.append(there)
        ends[s.sid] = (here, there)
    for s in sorted([s for s in secs if s.pred is not None], key=lambda s: s.installed):
        x = s.pred
        while x not in ends:
            x = by_id[x].pred
        ends[s.sid] = ends[x]
    _ENDS.clear()
    _ENDS[key] = ends
    return ends


def write_wo(path, rows):
    cols = ["wo_number", "feeder_id", "section_id", "fault_date", "completed_date", "cause_code", "repair_code",
            "crew_note"]
    us = lambda x: "%s/%s/%s" % (x[5:7], x[8:10], x[:4])      # the SAP user's date format (MM/DD/YYYY)
    with open(path, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=cols, lineterminator="\r\n")
        w.writeheader()
        for r in rows:
            w.writerow(dict(r, fault_date=us(r["fault_date"]), completed_date=us(r["completed_date"])))
    return len(rows)


# ---------------------------------------------------------------- GIS register
INS_CODE = {"PILC": "PILC", "XLPE": "XLPE", "TRXLPE": "TR-XLPE", "EPR": "EPR"}
CONDUCTOR = {"PILC": ["4/0 CU", "250 KCMIL CU", "350 KCMIL CU"], "XLPE": ["1/0 AL", "4/0 AL", "350 KCMIL AL"],
             "TRXLPE": ["1/0 AL", "4/0 AL", "750 KCMIL AL"], "EPR": ["1/0 AL", "4/0 AL", "500 KCMIL AL", "750 KCMIL AL"]}


# The GIS network was re-measured feeder by feeder in spring 2023 (field walk-downs and imagery), when the layer
# moved to the utility network model; every feature in service then got the new length and that edit date. The
# lengths before were carried over from the plat records.
REMEASURE = {"HR12": dt.date(2023, 3, 14), "HR14": dt.date(2023, 3, 28), "ML03": dt.date(2023, 4, 11),
             "ML07": dt.date(2023, 4, 25), "QH21": dt.date(2023, 5, 9), "QH22": dt.date(2023, 5, 23),
             "SP05": dt.date(2023, 6, 6), "SP09": dt.date(2023, 6, 20)}
PLAT_BIAS = {"HR12": 0.955, "HR14": 1.01, "ML03": 0.985, "ML07": 1.02, "QH21": 0.97, "QH22": 1.0, "SP05": 0.975,
             "SP09": 1.03}


def plat_length(s):
    """The section's length as the plat records carried it before the 2023 re-measurement."""
    if s.installed >= REMEASURE[s.feeder]:
        return s.length
    r = np.random.default_rng([W.SEED + 17, int(s.sid[2:])])
    return int(round(s.length * PLAT_BIAS[s.feeder] * float(r.normal(1.0, 0.035)) / 10) * 10)


def register_rows(secs, seed=W.SEED + 12):
    rng = np.random.default_rng(seed)
    ends = structures(secs)
    rows = []
    for s in sorted(secs, key=lambda s: s.sid):
        cond = CONDUCTOR[s.ins][int(rng.integers(len(CONDUCTOR[s.ins])))]
        fr, to = ends[s.sid]
        rm = REMEASURE[s.feeder] + dt.timedelta(days=int(rng.integers(0, 4)))
        if s.retired and s.retired < rm:
            edit = s.retired
        elif s.retired:
            edit = s.retired
        elif s.installed < rm:
            edit = rm
        else:
            edit = s.installed
        rows.append(dict(section_id=s.sid, feeder_id=current_feeder(s) if s.retired is None else feeder_at(s, s.retired),
                         from_structure=fr, to_structure=to, length_ft=s.length, conductor=cond,
                         insulation=INS_CODE[s.ins], voltage_class_kv=15 if s.ins != "PILC" else 15,
                         install_date=s.installed, install_method=s.method,
                         status="IS" if s.retired is None else "RET", retired_date=s.retired,
                         predecessor_section_id=s.pred or "", last_edit=edit))
    return rows


def write_register(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "UG_PRIMARY_SECTIONS"
    ws["A1"] = "%s  |  GIS export: underground primary cable sections, Tamsin Valley district" % W.COMPANY
    ws["A1"].font = Font(bold=True, size=11)
    ws["A2"] = "Data as of 2026-09-30 23:59 (nightly extract run 2026-10-01 02:15). Feeders HR12, HR14, ML03, ML07, QH21, QH22, SP05, SP09."
    ws["A2"].font = Font(italic=True, size=9)
    cols = ["section_id", "feeder_id", "from_structure", "to_structure", "length_ft", "conductor", "insulation",
            "voltage_class_kv", "install_date", "install_method", "status", "retired_date", "predecessor_section_id",
            "last_edit"]
    hdr = 4
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=hdr, column=j, value=c.upper())
        cell.font = Font(bold=True, size=9, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="44546A")
    for i, r in enumerate(rows, hdr + 1):
        for j, c in enumerate(cols, 1):
            v = r[c]
            cell = ws.cell(row=i, column=j, value=v if v != "" else None)
            if isinstance(v, dt.date):
                cell.number_format = "yyyy-mm-dd"
    widths = [12, 9, 13, 13, 9, 15, 10, 9, 11, 14, 7, 11, 14, 11]
    for j, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = wdt
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "A%d:%s%d" % (hdr, get_column_letter(len(cols)), hdr + len(rows))
    wb.save(path)
    return len(rows)


def write_codes(path):
    rows = [("CAUSE", "INS", "Cable insulation failure (primary cable faulted through its insulation)"),
            ("CAUSE", "SPL", "Splice failure (a cable joint failed)"),
            ("CAUSE", "TRM", "Termination failure (elbow, pothead or riser termination on the section failed)"),
            ("CAUSE", "TPD", "Third-party damage (dig-in or other damage by others)"),
            ("CAUSE", "OTH", "Other or no cable fault found (secondary, equipment, lightning, rodent, fire)"),
            ("REPAIR", "SPLICE", "Repaired in place with one or more splices"),
            ("REPAIR", "TERM", "New termination installed"),
            ("REPAIR", "REPL", "Section replaced: the failed cable is removed and its GIS section retired; the new cable is "
                               "added to the GIS as a new section"),
            ("STATUS", "IS", "In service"),
            ("STATUS", "RET", "Retired (removed from service; kept in the GIS with its retirement date)")]
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh, lineterminator="\r\n")
        w.writerow(["code_type", "code", "description"])
        w.writerows(rows)
    return len(rows)


# ---------------------------------------------------------------- grid: snapshots and restoration
class Snap:
    """Meters behind a device on a day, from the connectivity rows, and each meter's device chain."""

    def __init__(self, devices, meter_rows):
        self.devices = devices
        self.children = {}
        for d, v in devices.items():
            if v["parent"] is not None:
                self.children.setdefault(v["parent"], []).append(d)
        self.by_dev = {}
        for r in meter_rows:
            self.by_dev.setdefault(r["device_id"], []).append((r["meter_id"], r["eff_from"], r["eff_to"], r["feeder_id"]))
        self._desc, self._chain = {}, {}

    def desc(self, d):
        if d not in self._desc:
            out, stack = [], [d]
            while stack:
                x = stack.pop()
                out.append(x)
                stack.extend(self.children.get(x, []))
            self._desc[d] = out
        return self._desc[d]

    def chain(self, d):
        if d not in self._chain:
            out, x = [], d
            while x is not None:
                out.append(x)
                x = self.devices[x]["parent"]
            self._chain[d] = out
        return self._chain[d]

    def meters(self, d, day):
        out = []
        for dev in self.desc(d):
            for m, a, b, f in self.by_dev.get(dev, []):
                if a <= day and (b is None or day < b):
                    out.append((m, dev, f))
        return out


def restorations(snap, o):
    """[(meter, feeder, restore datetime, step index)] for one outage under the deepest-restoring-device rule."""
    day = o["start"].astimezone(W.LOCAL).date()
    depth = {r: len(snap.chain(r)) for r, _ in o["steps"]}
    out = []
    for m, dev, f in snap.meters(o["device"], day):
        ch = snap.chain(dev)
        best = None
        for i, (r, t) in enumerate(o["steps"]):
            if r in ch and (best is None or depth[r] > depth[o["steps"][best][0]] or
                            (depth[r] == depth[o["steps"][best][0]] and t > o["steps"][best][1])):
                best = i
        out.append((m, f, o["steps"][best][1], best))
    return out


def _clock(aware, system):
    if system == "OMS":
        return aware.astimezone(W.LOCAL).replace(tzinfo=None).strftime("%Y-%m-%d %H:%M:%S")
    return aware.astimezone(W.UTC).replace(tzinfo=None).strftime("%Y-%m-%d %H:%M:%S")


def oms_tables(snap, outs, ship_from=W.OMS_START):
    """The outage export rows (one per outage and affected feeder) and the step rows, for the outages starting on or
    after ship_from (local)."""
    oms, steps = [], []
    seq_l, seq_a = {}, [4100]
    for o in outs:
        if o["start"].astimezone(W.LOCAL).date() < ship_from:
            continue
        system = "OMS" if o["start"] < W.ADMS_CUTOVER else "ADMS"
        if system == "OMS":
            key = o["start"].astimezone(W.LOCAL).strftime("%y%m%d")
            seq_l[key] = seq_l.get(key, 0) + 1
            oid = "OMS-%s-%03d" % (key, seq_l[key])
        else:
            seq_a[0] += 1 + (seq_a[0] * 7) % 5
            oid = "INC%s%05d" % (o["start"].astimezone(W.UTC).strftime("%y"), seq_a[0])
        o["outage_id"], o["system"] = oid, system
        res = restorations(snap, o)
        o["res"] = res
        n_total = len(res)
        cm_total = sum((t - o["start"]).total_seconds() / 60 for _, _, t, _ in res)
        dtype = snap.devices[o["device"]]["type"]
        for f in o["feeders"]:
            oms.append(dict(outage_id=oid, feeder_id=f, device_id=o["device"], device_type=dtype,
                            start_ts=_clock(o["start"], system), restore_ts=_clock(max(t for _, t in o["steps"]), system),
                            customers_out=n_total, cust_minutes=round(cm_total, 1),
                            cause_code=o["cause"] if system == "OMS" else ADMS_CAUSE.get(o["cause"], o["cause"].title()),
                            source_system=system))
        for i, (r, t) in enumerate(o["steps"]):
            steps.append(dict(outage_id=oid, step_no=i + 1, device_id=o["device"], restoring_device_id=r,
                              step_restore_ts=_clock(t, system), customers_restored=sum(1 for x in res if x[3] == i),
                              source_system=system))
    return oms, steps


def write_csv(path, rows, cols):
    with open(path, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=cols, lineterminator="\r\n")
        w.writeheader()
        for r in rows:
            w.writerow({c: ("" if r.get(c) is None else r[c]) for c in cols})
    return len(rows)


# ADMS records the dispatcher's cause in its own pick list; the legacy system's codes came across as they were
ADMS_CAUSE = {"UG FAULT": "Underground Cable Fault", "EQUIPMENT": "Equipment Failure", "ANIMAL": "Animal Contact",
              "VEHICLE": "Vehicle Accident", "UNKNOWN": "Unknown", "DIG-IN": "Dig-In/Third Party",
              "SUBSTATION": "Substation Equipment", "STORM": "Weather - Wind/Storm", "WEATHER": "Weather - Wind/Storm",
              "TREE": "Vegetation"}
OMS_COLS = ["outage_id", "feeder_id", "device_id", "device_type", "start_ts", "restore_ts", "customers_out",
            "cust_minutes", "cause_code", "source_system"]
STEP_COLS = ["outage_id", "step_no", "device_id", "restoring_device_id", "step_restore_ts", "customers_restored",
             "source_system"]


def write_connectivity(path, t):
    rows = sorted(t.meter_rows, key=lambda r: (r["feeder_id"], r["transformer_id"], r["meter_id"], r["eff_from"]))
    df = pd.DataFrame({
        "meter_id": [r["meter_id"] for r in rows],
        "meter_serial": [r["serial"] for r in rows],
        "transformer_id": [r["transformer_id"] for r in rows],
        "device_id": [r["device_id"] for r in rows],
        "feeder_id": [r["feeder_id"] for r in rows],
        "eff_from": pd.to_datetime([r["eff_from"] for r in rows]).date,
        "eff_to": [r["eff_to"] for r in rows],
    })
    import pyarrow as pa
    import pyarrow.parquet as pq
    table = pa.Table.from_pandas(df, preserve_index=False)
    table = table.replace_schema_metadata({})
    pq.write_table(table, path, compression="snappy", write_statistics=False, use_dictionary=True)
    return len(rows)


def write_devices(path, t):
    subs = []
    for sub, name, feeders in W.SUBSTATIONS:
        subs.append({"substation_id": sub, "name": "%s Substation" % name,
                     "banks": [{"bank_id": "%s-TB1" % sub, "kva": 33600 if sub != "QHL" else 28000,
                                "feeders": [{"feeder_id": f, "circuit": W.CIRCUIT_NO[f], "name": W.NAME[f],
                                             "breaker": "CB-%s" % f, "kv": 13.2} for f in feeders]}]})
    devs = []
    for d in sorted(t.devices, key=lambda x: (t.devices[x]["sub"], t.devices[x]["feeder"] or "", x)):
        v = t.devices[d]
        devs.append({"device_id": d, "device_type": v["type"], "feeder_id": v["feeder"],
                     "parent_device_id": v["parent"],
                     "in_service_from": v["in_from"].isoformat() if v["in_from"] else None,
                     "in_service_to": v["in_to"].isoformat() if v["in_to"] else None})
    doc = {"export": "Distribution network model, protective device hierarchy", "district": "Tamsin Valley",
           "model_date": "2026-09-30", "substations": subs, "devices": devs}
    with open(path, "w") as fh:
        json.dump(doc, fh, indent=1)
        fh.write("\n")
    return len(devs)


STREETS = ["Harrow Rd", "Mill Ln", "Quarry Hill Rd", "Sutter Park Dr", "Linden St", "Kessler Ave", "Tamsin Valley Rd",
           "Orchard Way", "Brookside Dr", "Fairview Ave", "Spruce Ct", "Ridgeline Dr", "Canal St", "Juniata Ave",
           "Penn Ave", "Schuylkill Rd", "Furnace Rd", "Old Forge Ln", "Wren Ct", "Birchwood Ln", "Hollow Rd",
           "Chestnut St", "Kittatinny Dr", "Mountain View Ter", "Laurel Run Rd", "Station Rd", "Beaver Ln",
           "Carriage Hill Dr", "Meadowbrook Ln", "Summit Ave"]


def cis_rows(t, seed=W.SEED + 13):
    rng = np.random.default_rng(seed)
    first = {}
    for r in sorted(t.meter_rows, key=lambda r: r["eff_from"]):
        first.setdefault(r["meter_id"], r)
    rows = []
    acct = 30_118_402
    live = {r["meter_id"] for r in t.meter_rows if r["eff_from"] <= W.CUT and (r["eff_to"] is None or W.CUT < r["eff_to"])}
    tx_street = {}
    for m in sorted(live, key=lambda m: (first[m]["transformer_id"], m)):
        r0 = first[m]
        acct += int(rng.integers(3, 60))
        if r0["eff_from"] == dt.date(2014, 11, 9):
            setup = dt.date(int(rng.integers(1988, 2014)), int(rng.integers(1, 13)), int(rng.integers(1, 29)))
        else:
            setup = r0["eff_from"] - dt.timedelta(days=int(rng.integers(3, 40)))
        st = tx_street.setdefault(r0["transformer_id"], STREETS[int(rng.integers(len(STREETS)))])
        rc = "RS" if rng.random() < 0.88 else ("GS" if rng.random() < 0.8 else "GS-L")
        rows.append(dict(account_no=str(acct), meter_id=m, service_address="%d %s" % (int(rng.integers(2, 2400)), st),
                         rate_class=rc, account_setup_date=setup.isoformat(), feeder_label=r0["feeder_id"],
                         premise_type="RES" if rc == "RS" else "COM"))
    rows.sort(key=lambda r: r["account_no"])
    return rows


CIS_COLS = ["account_no", "meter_id", "service_address", "rate_class", "account_setup_date", "feeder_label",
            "premise_type"]


def write_cis(path, rows):
    """The customer system's extract job: upper-case field names, dates as YYYYMMDD, addresses as the billing system
    holds them (upper case, fixed-width field)."""
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh, lineterminator="\r\n")
        w.writerow([c.upper() for c in CIS_COLS])
        for r in rows:
            w.writerow([r["account_no"], r["meter_id"], r["service_address"].upper().ljust(30), r["rate_class"],
                        r["account_setup_date"].replace("-", ""), r["feeder_label"], r["premise_type"]])
    return len(rows)


def write_med(path, rows):
    cols = ["reliability_year", "med_date", "tmed_saidi_min", "event"]
    return write_csv(path, rows, cols)


VAULT_FINDINGS = ["no defects noted", "no defects noted", "no defects noted", "minor standing water, pumped",
                  "cable racking corroded", "lid frame cracked", "debris in vault, cleaned", "splice tape weathered", "silt 2½ in, cleaned",
                  "fireproofing tape missing on one cable", "ground wire loose, retightened", "silt 4 in, cleaned", "water 18½ in below racking, pumped",
                  "water to racking level, pumped", "ladder corroded", "neutral corrosion visible on one cable",
                  "duct seal missing at one duct", "oil sheen on water, reported to environmental"]


# the inspection crews over the years (initials as each typed them; two-person crews enter both)
INSPECTORS = [("JRK", 2022, 2026, 5), ("MDT", 2022, 2024, 4), ("LAP", 2022, 2026, 3), ("BSH", 2022, 2023, 2),
              ("TNW", 2023, 2026, 4), ("jrk", 2022, 2026, 1), ("R. Okafor", 2024, 2026, 3), ("DVS", 2025, 2026, 3),
              ("JRK/TNW", 2023, 2026, 1), ("CONTR-KLINE", 2022, 2022, 2), ("MDT ", 2022, 2024, 1)]


def vault_rows(secs, seed=W.SEED + 14):
    rng = np.random.default_rng(seed)
    ends = structures(secs)
    structs = sorted({(x, current_feeder(s)) for s in secs if s.method == "CONDUIT" and s.retired is None
                      for x in ends[s.sid] if x.startswith("MH-")})
    seen, uniq = set(), []
    for sid, f in structs:
        if sid not in seen:
            seen.add(sid)
            uniq.append((sid, f))
    rows = []
    for sid, f in uniq:
        n = int(rng.integers(1, 4))
        first = dt.date(2022, 1, 3) + dt.timedelta(days=int(rng.integers(0, 600)))
        for i in range(n):
            d = first + dt.timedelta(days=int(i * rng.integers(330, 420)))
            if d > dt.date(2026, 9, 25):
                break
            wl = str(rng.choice(["DRY", "DAMP", "STANDING", "STANDING", "DAMP", "DRY", "SUBMERGED"],
                                p=[0.3, 0.25, 0.12, 0.12, 0.1, 0.08, 0.03]))
            cond = int(rng.choice([1, 2, 3, 4, 5], p=[0.05, 0.15, 0.42, 0.3, 0.08]))
            find = VAULT_FINDINGS[int(rng.integers(len(VAULT_FINDINGS)))]
            crews = [(c, wt) for c, y0, y1, wt in INSPECTORS if y0 <= d.year <= y1]
            wts = np.array([wt for _, wt in crews], dtype=float)
            who = crews[int(rng.choice(len(crews), p=wts / wts.sum()))][0]
            rows.append(dict(structure_id=sid, structure_type="MANHOLE" if int(sid[3:]) % 3 else "VAULT", feeder_id=f,
                             inspection_date=d, inspector=who, water_level=wl, condition_rating=cond,
                             findings=find, follow_up="Y" if cond <= 2 or wl == "SUBMERGED" else "N"))
    rows.sort(key=lambda r: (r["inspection_date"], r["structure_id"]))
    return rows


VAULT_COLS = ["structure_id", "structure_type", "feeder_id", "inspection_date", "inspector", "water_level",
              "condition_rating", "findings", "follow_up"]
# the inspections database is an Access application; its export is Windows-1252 with its own column captions
VAULT_HEAD = ["Structure ID", "Structure Type", "Feeder", "Insp Date", "Inspector", "Water Level", "Condition (1-5)",
              "Findings", "Follow-up WO"]


def write_vaults(path, rows):
    with open(path, "w", newline="", encoding="cp1252") as fh:
        w = csv.writer(fh, lineterminator="\r\n")
        w.writerow(VAULT_HEAD)
        for r in rows:
            d = r["inspection_date"]
            w.writerow([r["structure_id"], r["structure_type"], r["feeder_id"], "%d/%d/%d" % (d.month, d.day, d.year),
                        r["inspector"], r["water_level"], r["condition_rating"], r["findings"], r["follow_up"]])
    return len(rows)


RATING = {"HR12": 530, "HR14": 600, "ML03": 600, "ML07": 600, "QH21": 530, "QH22": 600, "SP05": 530, "SP09": 600}
LOAD = {"HR12": 0.62, "HR14": 0.66, "ML03": 0.7, "ML07": 0.64, "QH21": 0.48, "QH22": 0.55, "SP05": 0.8, "SP09": 0.52}


def write_loading(path, seed=W.SEED + 15):
    rng = np.random.default_rng(seed)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Monthly peaks"
    ws["A1"] = "Feeder monthly peak loading (SCADA feeder breaker amps), Tamsin Valley district"
    ws["A1"].font = Font(bold=True)
    ws["A2"] = "Source: SCADA historian, monthly maximum of 15-minute average phase current. Rating = feeder getaway cable emergency rating."
    ws["A2"].font = Font(italic=True, size=9)
    cols = ["Feeder", "Month", "Peak amps", "Rating amps", "Peak % of rating"]
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=4, column=j, value=c)
        cell.font = Font(bold=True)
    r = 5
    for f in W.FEEDERS:
        for y in range(2021, 2027):
            for m in range(1, 13):
                if (y, m) > (2026, 9):
                    break
                season = 1.0 + 0.32 * max(0.0, np.cos((m - 7.3) / 12 * 2 * np.pi)) + 0.08 * (m in (1, 2, 12))
                hot = {2022: 1.06, 2023: 1.04}.get(y, 1.0) if m in (6, 7, 8) else 1.0
                grow = 1 + 0.012 * (y - 2021) + (0.09 if f == "SP05" and (y, m) >= (2021, 5) and (y, m) < (2025, 4) else 0)
                grow += 0.11 if f == "ML07" and (y, m) >= (2026, 6) else 0
                amps = RATING[f] * LOAD[f] * season * hot * grow * float(rng.normal(1, 0.025))
                ws.cell(row=r, column=1, value=f)
                c = ws.cell(row=r, column=2, value=dt.date(y, m, 1))
                c.number_format = "mmm-yy"
                ws.cell(row=r, column=3, value=int(round(amps)))
                ws.cell(row=r, column=4, value=RATING[f])
                c = ws.cell(row=r, column=5, value=round(int(round(amps)) / RATING[f], 4))
                c.number_format = "0%"
                r += 1
    for j, wdt in enumerate([9, 9, 10, 11, 15], 1):
        ws.column_dimensions[get_column_letter(j)].width = wdt
    ws.freeze_panes = "A5"
    wb.save(path)
    return r - 5
