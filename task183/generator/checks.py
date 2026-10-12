"""Assertions over the built pack and the goldens: the call, every rung by name with its margin, the 16-cell grid
cell by cell and the partial cells, the semantic fork rows SF1 to SF9, the six readability tests, the clean-data
and lens tests, every ask figure (bins, major event day margins) and every designed wrong stop, the separation of
the ask devices from the main call's rows, the distractors, the report and the published major event days, the
pack gates, dates, vocabulary and containers. Everything is computed from target/ and golden/ (models.py, asks.py),
plus the simulated world where an assertion is about the world itself."""
import datetime as dt
import json
import math
import os
import re
import shutil
import sys
import tempfile
import zipfile

import numpy as np
import openpyxl
import pandas as pd
import pdfplumber

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import asks as A        # noqa: E402
import build as B       # noqa: E402
import cable as C       # noqa: E402
import deliver as DL    # noqa: E402
import load as L        # noqa: E402
import models as M      # noqa: E402
import records as R     # noqa: E402
import world as W       # noqa: E402

TASK = os.path.dirname(HERE)
REPO = os.path.dirname(TASK)
N = [0]
LOG = []


FAILS = []


def ok(cond, msg):
    N[0] += 1
    if not cond:
        if os.environ.get("T183_SOFT"):
            FAILS.append(msg)
            print("FAIL:", msg)
            return
        raise AssertionError(msg)


def lead(fc):
    o, l = M.rank(fc)
    return o[0], float(l), o.index("QH21") + 1, o


def note(name, fc):
    w, l, q, o = lead(fc)
    LOG.append("%-58s %s by %.2f, Quarry Hill 21 #%d | %s" % (name, w, l, q, " ".join("%s %.1f" % (f, fc[f]) for f in o)))
    return w, l, q, o


def bin_dist(x, dp):
    s = x * 10 ** dp
    return abs((s - math.floor(s)) - 0.5)


def run(target, golden, world=None, meta_path=None, quiet=False):
    wo, reg = L.work_orders(target), L.register(target)
    oms, steps, conn, dv, cis = L.oms(target), L.steps(target), L.connectivity(target), L.devices(target), L.cis(target)
    meds_pub = L.meds(target)

    # ------------------------------------------------------------ the ladder, by name
    r0 = note("R0 five-year average, retired sections included", M.history(wo, reg, pool="all"))
    r1 = note("R1 five-year average on sections in service (default)", M.history(wo, reg))
    r2 = note("R2 rates by insulation and age band on the stock", M.attributes(wo, reg))
    r3 = note("R3 rates by failures already on the section", M.kmodel(wo, reg)[0])
    ok(r0[0] == "HR12" and r0[1] >= 1.4, "R0 names Harrow Road 12 by 1.4 or more: %s" % (r0[:2],))
    ok(r1[0] == "ML07" and r1[1] >= 1.12, "R1 names Mill Lane 7 by 1.12 or more: %s" % (r1[:2],))
    ok(r2[0] == "SP05" and r2[1] >= 1.12, "R2 names Sutter Park 5 by 1.12 or more: %s" % (r2[:2],))
    ok(r3[0] == "QH21" and r3[1] >= 1.45, "R3 names Quarry Hill 21 by 1.45 or more: %s" % (r3[:2],))
    # position rule: fifth or fourth on the default path, never first or second on an intermediate rung
    ok(r1[2] in (4, 5), "Quarry Hill 21 is 4th or 5th on the default path: %d" % r1[2])
    ok(r0[2] >= 4 and r2[2] >= 3, "Quarry Hill 21 is not first or second on R0 or R2: %d, %d" % (r0[2], r2[2]))
    fc1 = M.history(wo, reg)
    above = sorted(fc1.values(), reverse=True)[r1[2] - 2]
    ok(above / fc1["QH21"] >= 1.05, "the feeder above Quarry Hill 21 on the default path leads it by 5 per cent")
    ok(len({r0[0], r1[0], r2[0], r3[0]}) == 4, "the four rungs name four different feeders")
    # the decisive edge: Quarry Hill 21 is the only feeder with in-service sections at two or more failures
    _, fails = M.prep(wo, reg)
    st = M.stock(reg, fails)
    twice = st[st["k"] >= 2].groupby("feeder_id").size().to_dict()
    ok(set(twice) == {"QH21"} and twice["QH21"] >= 12, "only Quarry Hill 21 carries twice-failed sections: %s" % twice)
    r = M.kmodel(wo, reg)[1]
    ok(r[2] / r[1] >= 4 and r[2] >= 0.25 and r[3] >= r[2], "the record's rate after two failures dwarfs the rate "
                                                             "after one: %s" % r.round(4).to_dict())
    ok(r[2] >= 10 * r[1], "after two failures a section fails ten times as often as after one: %s" % r.round(4).to_dict())

    # ------------------------------------------------------------ R3 variants (rigorous routes all land)
    variants = {
        "k from 2012": M.kmodel(wo, reg)[0],
        "k from 2017": M.kmodel(wo, reg, start=dt.date(2017, 1, 1))[0],
        "k from 2019": M.kmodel(wo, reg, start=dt.date(2019, 1, 1))[0],
        "k classes 0,1,2+": M.kmodel(wo, reg, top=2)[0],
        "k x insulation x age band (10-year)": M.k_attributes(wo, reg),
        "k x insulation x age band (5-year)": M.k_attributes(wo, reg, width=5),
        "Poisson GLM, k + insulation + age band": M.poisson_glm(wo, reg),
        "Poisson GLM, k + insulation + age band, from 2017": M.poisson_glm(wo, reg, start=dt.date(2017, 1, 1)),
        "Poisson GLM, k only": M.poisson_glm(wo, reg, attrs=False),
        "Poisson GLM, k + insulation + age band + feeder": M.poisson_glm(wo, reg, feeder_effect=True),
    }
    floors = {}
    for name, fc in variants.items():
        w, l, q, o = note("R3 variant: " + name, fc)
        ok(w == "QH21" and l >= floors.get(name, 1.25), "R3 variant %s names Quarry Hill 21 (%s %.2f)" % (name, w, l))

    # ------------------------------------------------------------ the 16-cell grid
    grid = {}
    for pool in ("in_service", "all"):
        for causes, ctag in ((M.CABLE, "cable"), (None, "all work orders")):
            for model in ("history", "attributes", "k", "k+attr"):
                if model == "history":
                    fc = M.history(wo, reg, causes, pool)
                elif model == "attributes":
                    fc = M.attributes(wo, reg, causes, pool)
                elif model == "k":
                    fc = M.kmodel(wo, reg, causes, pool)[0]
                else:
                    fc = M.k_attributes(wo, reg, causes, pool)
                grid[(pool, ctag, model)] = note("grid %s / %s / %s" % (pool, ctag, model), fc)
    for (pool, ctag, model), (w, l, q, o) in grid.items():
        if pool == "in_service" and model in ("k", "k+attr"):
            ok(w == "QH21" and l >= 1.1, "grid cell %s/%s/%s names Quarry Hill 21 by 1.1 (%s %.2f)" % (
                pool, ctag, model, w, l))
        elif pool == "in_service":
            ok(w in ("ML07", "SP05") and l >= 1.05, "grid cell %s/%s/%s names %s" % (pool, ctag, model, w))
        else:
            # violates the prompt's "sections in service at the end of September 2026": retired cable is counted
            ok(w == "HR12" and l >= 1.2, "grid cell %s/%s/%s (retired cable counted) names Harrow Road 12" % (
                pool, ctag, model))
    ok(sum(1 for v in grid.values() if v[0] == "QH21") == 4, "exactly the four in-service failure-count cells name "
                                                               "Quarry Hill 21")

    # ------------------------------------------------------------ partial cells
    pw = note("partial: failure count inside the five-year window only", M.kmodel_window_only(wo, reg))
    ok(pw[0] == "QH21", "the count taken inside the window still names Quarry Hill 21")
    pi = note("partial: history passed to replacement sections", M.kmodel_inherited(wo, reg))
    ok(pi[0] == "QH21" and pi[1] >= 1.2, "history passed to replacements still names Quarry Hill 21")
    pt = note("partial: straight line through yearly counts", M.trend(wo, reg))
    ok(pt[0] == "SP05", "a trend on yearly counts names Sutter Park 5: %s" % pt[0])
    py = note("partial: last twelve months' failures", M.last_year(wo, reg))
    ok(py[0] == "SP05" and py[1] >= 1.05, "the one-period count names Sutter Park 5: %s" % (py[:2],))
    fe = {"Poisson GLM, k + feeder": M.poisson_glm(wo, reg, feeder_effect=True, attrs=False),
          "Poisson GLM, k + feeder, from October 2021": M.poisson_glm(wo, reg, feeder_effect=True, attrs=False,
                                                                       start=dt.date(2021, 10, 1)),
          "Poisson GLM, k + attributes + feeder, from October 2021": M.poisson_glm(wo, reg, feeder_effect=True,
                                                                                    start=dt.date(2021, 10, 1))}
    fe_res = {k: note("partial (feeder effect): " + k, v) for k, v in fe.items()}
    # a feeder effect without the cable's attributes carries Harrow Road 12's removed paper cable into its new EPR:
    # the record refutes it (the feeder's in-service cable has almost no failures)
    ok(fe_res["Poisson GLM, k + feeder"][0] == "HR12" and
       fe["Poisson GLM, k + feeder"]["HR12"] >= 10 * max(fc1["HR12"], 0.2),
       "a feeder effect alone carries Harrow Road 12's removed PILC into its new cable (%.1f against %.1f a year)" % (
           fe["Poisson GLM, k + feeder"]["HR12"], fc1["HR12"]))
    # on the five-year window alone the failure count is not identified apart from the feeder: only Quarry Hill 21
    # has section years at two or more failures since October 2021
    ex21 = M.exposures(reg, fails, start=dt.date(2021, 10, 1))
    k2 = ex21[ex21["k"] >= 2]
    ok(k2.loc[k2["feeder_id"] == "QH21", "years"].sum() >= 0.9 * k2["years"].sum(), "since October 2021 nine tenths "
       "of the section years at two or more failures are Quarry Hill 21's")
    ly = M.last_year(wo, reg)
    ml7 = wo[(wo["feeder_id"] == "ML07") & wo["cause_code"].isin(M.CABLE)]["fault_date"].map(lambda d: d.year)
    yc = [int((ml7 == y).sum()) for y in (2022, 2023, 2024, 2025)]
    ok(yc[0] > yc[1] > yc[2] > yc[3], "Mill Lane 7's cable failures fell each year 2022 to 2025: %s" % yc)
    ok(ly["QH21"] >= 2 * ly["ML07"], "Quarry Hill 21's last twelve months are twice Mill Lane 7's")

    # ------------------------------------------------------------ semantic fork rows
    sf = {
        "SF1 every cable work order, dig-ins included": M.kmodel(wo, reg, causes=None)[0],
        "SF2 the twelve months after the cut": M.kmodel(wo, reg)[0],
        "SF3 projected to 1 January 2027": M.kmodel_projected(wo, reg),
        "SF4 feeder from the work orders": M.by_wo_feeder(wo, reg),
        "SF5 failures per mile": M.per_mile(M.kmodel(wo, reg)[0], reg),
        "SF6 completion date for the failure date": M.by_completion(wo, reg),
        "SF7 history passed through the predecessor pointer": M.kmodel_inherited(wo, reg),
        "SF8 hazard window from 2019": M.kmodel(wo, reg, start=dt.date(2019, 1, 1))[0],
    }
    for name, fc in sf.items():
        w, l, q, o = note(name, fc)
        if name.startswith("SF7"):
            # a replacement section is a new section in the code list and the register; handing it the removed
            # cable's failures still names Quarry Hill 21, and the replacements' own record shows why it is wrong
            codes = open(os.path.join(target, R.CODES)).read()
            ok("the new cable is added to the GIS as a new section" in codes, "SF7: the code list files a replacement "
                                                                              "as a new section")
            ok(w == "QH21" and l >= 1.2, "SF7 gives Quarry Hill 21 (%s %.2f)" % (w, l))
            pred = dict(zip(reg["section_id"], reg["predecessor_section_id"]))
            yrs, ev = 0.0, 0
            for row in reg[reg["predecessor_section_id"] != ""].itertuples(index=False):
                x, inherited = pred.get(row.section_id), 0
                while x:
                    inherited += len(fails.get(x, []))
                    x = pred.get(x)
                if inherited >= 2:
                    end = row.retired_date or W.CUT
                    yrs += (end - row.install_date).days / 365.25
                    ev += len(fails.get(row.section_id, []))
            ok(yrs > 100 and ev / yrs < 0.05, "SF7: replacements of twice-failed cable fail at under 5 per cent a "
                                              "year (%d in %.0f section years)" % (ev, yrs))
        else:
            ok(w == "QH21" and l >= 1.15, "%s gives Quarry Hill 21 (%s %.2f)" % (name, w, l))
    # SF3: no replacement or install recorded after the cut
    ok(all(d is None or d <= W.CUT for d in reg["retired_date"]) and all(d <= W.CUT for d in reg["install_date"]),
       "SF3: no GIS change dated after 30 September 2026")
    # SF4: the moved lateral carries no twice-failed section
    lat = wo[(wo["feeder_id"] == "ML03")]
    moved = set(reg.loc[(reg["feeder_id"] == "ML07") & reg["section_id"].isin(set(lat["section_id"])), "section_id"])
    ok(len(moved) > 0 and all(len(fails.get(s, [])) < 2 for s in moved), "SF4: the moved lateral has failures and "
                                                                          "no twice-failed section")
    # SF6: no failure completed in a later year than its fault would move a section's count past the cut
    ok(all(c >= f for f, c in zip(wo["fault_date"], wo["completed_date"])), "SF6: completion never before the fault")
    # SF9: Sutter Park 5's recent lead is the Kessler Avenue loop's overload (the EPR laid 2006 to 2010), which the
    # report (6.3) and the extract notes file as ended by the loop's second feed on 15 September 2026, so carrying
    # the loop's overload years into 2027 is the reading those sentences rule out
    loop = set(reg.loc[(reg["feeder_id"] == "SP05") & (reg["insulation"] == "EPR") &
                       reg["install_date"].map(lambda d: 2006 <= d.year <= 2010), "section_id"])
    lf = sorted(d for s in loop for d in fails.get(s, []))
    relief = dt.date(2026, 9, 15)
    ok(len(loop) == 72 and lf and lf[-1] < relief,
       "SF9: the last of the 72 loop sections' failures (%s) precedes the loop's second feed" % (lf[-1] if lf else None))
    sp_cable = wo[(wo["feeder_id"] == "SP05") & wo["cause_code"].isin(M.CABLE)]
    n25 = int((sp_cable["fault_date"].map(lambda d: d.year) == 2025).sum())
    l25 = sum(1 for d in lf if d.year == 2025)
    ok(2 * l25 > n25, "SF9: most of Sutter Park 5's 2025 cable failures are on the loop (%d of %d)" % (l25, n25))
    notes_txt = open(os.path.join(target, "asset_analytics_extract_notes_2026-10.md")).read()
    ok(notes_txt.count("second feed to the Kessler Avenue loop was energised on 15 Sep 2026") == 1,
       "SF9: the extract notes file the loop's second feed once")
    with pdfplumber.open(os.path.join(target, "2025_annual_reliability_report_sec6_ug_cable.pdf")) as pdf:
        rep_txt = " ".join((pg.extract_text() or "") for pg in pdf.pages).replace("\n", " ")
    ok("Kessler Avenue loop" in rep_txt and "above the loop's rating" in rep_txt,
       "SF9: the 2025 report (6.3) files the loop's overload since the 2024 expansion")
    hr14 = wo[(wo["feeder_id"] == "HR14") & wo["cause_code"].isin(M.CABLE)]["fault_date"].map(lambda d: d.year)
    before = np.mean([int((hr14 == y).sum()) for y in (2017, 2018, 2019)])
    after = [int((hr14 == y).sum()) for y in (2025, 2026)]
    ok("ended with the elbows' change-out in 2024" in rep_txt and max(after) <= before + 1,
       "SF9: the report (6.4) files Harrow Road 14's elbow run as ended, and its 2025 and 2026 failures (%s) are back at "
       "its 2017 to 2019 level (%.1f a year)" % (after, before))
    # with the loop's overload years taken out of its record (its 2024 to 2026 failures), every recent-window count and
    # the failure-count model name Quarry Hill 21; with them left in, the recent windows name Sutter Park 5 (the rung)
    over = {(s, d) for s in loop for d in fails.get(s, []) if d.year >= 2024}
    insvc = set(reg.loc[reg["status"] == "IS", "section_id"])
    for lo, label in ((dt.date(2025, 10, 1), "last twelve months"), (dt.date(2024, 10, 1), "last two years")):
        cnt_all, cnt_net = {}, {}
        for f in W.FEEDERS:
            secs_f = set(reg.loc[(reg["feeder_id"] == f) & reg["section_id"].isin(insvc), "section_id"])
            ev = [(s, d) for s in secs_f for d in fails.get(s, []) if lo <= d <= W.CUT]
            cnt_all[f] = len(ev)
            cnt_net[f] = len([e for e in ev if e not in over])
        top_all = max(cnt_all, key=cnt_all.get)
        top_net = max(cnt_net, key=cnt_net.get)
        ok(top_all == "SP05" and top_net == "QH21",
           "SF9: the %s name Sutter Park 5 with the loop's overload kept (%d) and Quarry Hill 21 without it (%d)"
           % (label, cnt_all["SP05"], cnt_net["QH21"]))

    # ------------------------------------------------------------ readability tests (the decisive rung silent)
    texts = []
    for f in os.listdir(target):
        p = os.path.join(target, f)
        if f.endswith(".md") or f.endswith(".csv") and os.path.getsize(p) < 2_000_000:
            raw = open(p, "rb").read()
            try:
                texts.append((f, raw.decode("utf-8")))
            except UnicodeDecodeError:
                texts.append((f, raw.decode("cp1252")))
        elif f.endswith(".pdf"):
            with pdfplumber.open(p) as pdf:
                texts.append((f, "\n".join(pg.extract_text() or "" for pg in pdf.pages)))
        elif f.endswith(".json"):
            texts.append((f, open(p).read()))
    banned = re.compile(r"repeat|recurr|second failure|2nd failure|failed twice|twice[- ]failed|failure count|"
                        r"prior failure|previous failure|failures so far|bad actor|chronic", re.I)
    hits = [(f, m.group(0)) for f, t in texts for m in banned.finditer(t)]
    ok(not hits, "test 1: no shipped sentence names repeat failure or a section's failure count: %s" % hits[:5])
    cols = set(wo.columns) | set(reg.columns)
    ok(not any(re.search(r"count|prior|repeat|n_fail|num_fail|history", c) for c in cols),
       "test 2: no column carries a section's failure count: %s" % sorted(cols))
    ok(r2[0] != "QH21", "test 3: the reflex (rates by visible attributes) does not name Quarry Hill 21")
    rep = report_check(target, wo, reg, conn)
    ok(rep, "test 4: the annual report reproduces from the work orders, the register and the connectivity model")
    ok(wo["wo_number"].is_unique and not wo.duplicated().any(), "test 5: work order numbers unique, no exact duplicates")
    ok(set(wo["section_id"]) <= set(reg["section_id"]), "test 5: every work order's section is in the register")
    ok(reg["section_id"].is_unique, "test 5: section ids unique in the register")
    preds = set(x for x in reg["predecessor_section_id"] if x)
    ok(preds <= set(reg.loc[reg["status"] == "RET", "section_id"]), "test 5: every predecessor is a retired section")
    ok(all(lead_[0] != "QH21" for lead_ in (r1, py if py[0] != "QH21" else ("X",), pt)),
       "test 6: no one-step arithmetic on counts names Quarry Hill 21")

    # ------------------------------------------------------------ Gate G: clean data, lens swap, identity
    clean = wo[wo["cause_code"].isin(M.CABLE)]
    ok(lead(M.kmodel(clean, reg)[0])[0] == "QH21" and lead(M.history(clean, reg))[0] == "ML07",
       "clean-data test: answer and naive answer unchanged on cause-filtered records, and still differ")
    ok(r1[0] != r3[0], "lens swap: the trailing five years and calendar 2027 name different feeders")

    # ------------------------------------------------------------ the world: the record's generating rates
    if world is not None:
        tr = C.truth_2027(world["secs"])
        w_, l_, _, _ = note("world: expected 2027 failures under the generating rates", tr)
        ok(w_ == "QH21" and l_ >= 1.3, "the world's own 2027 expectation names Quarry Hill 21 by 1.3")
        nf = {f: 0 for f in W.FEEDERS}
        for s in world["secs"]:
            if s.in_service(W.CUT) and s.k >= 2:
                nf[R.current_feeder(s)] += 1
        ok(nf == {f: twice.get(f, 0) for f in W.FEEDERS}, "the register and work orders reproduce the world's "
                                                          "twice-failed sections")

    # ------------------------------------------------------------ work order and register hygiene
    ok(not (wo["fault_date"] == W.CUT).any(), "no fault dated on the 30 September 2026 boundary")
    ok(wo["fault_date"].max() <= W.CUT and wo["completed_date"].max() <= W.CUT, "every work order closed by the cut")
    ok(wo["fault_date"].min() >= W.ARCHIVE_START, "work orders start with the 2012 archive")
    ok(set(wo["cause_code"]) == {"INS", "SPL", "TRM", "TPD", "OTH"}, "cause codes as the code list")
    ok(set(wo["repair_code"]) <= {"SPLICE", "TERM", "REPL"}, "repair codes as the code list")
    repl = wo[wo["repair_code"] == "REPL"]
    ok(set(repl["section_id"]) <= set(reg.loc[reg["status"] == "RET", "section_id"]), "every REPL order's section "
                                                                                       "is retired")
    hwy_ids = set()
    if world is not None:
        hwy_ids = {s.sid for s in world["secs"] if s.hwy}
        ok(all(s in set(reg.loc[reg["status"] == "IS", "section_id"]) for s in hwy_ids), "the State Route 443 run is "
                                                                                          "all in service")

    # ------------------------------------------------------------ the asks
    g = A.golden(oms, steps, conn, dv)
    figs = json.load(open(os.path.join(HERE, "figures.json")))
    for f in W.FEEDERS:
        ok("%.2f" % g["saifi"][f] == figs["saifi_ry2026"][f], "figures.json SAIFI %s" % f)
        ok("%.1f" % g["saidi"][f] == figs["saidi_ry2026_min"][f], "figures.json SAIDI %s" % f)
        ok(g["cemi4"][f] == figs["customers_4plus_ry2026"][f], "figures.json 4+ count %s" % f)
        ok(bin_dist(g["saifi"][f], 2) >= 0.2, "SAIFI %s mid-bin: %.5f" % (f, g["saifi"][f]))
        ok(bin_dist(g["saidi"][f], 1) >= 0.2, "SAIDI %s mid-bin: %.4f" % (f, g["saidi"][f]))
        ok(g["cemi4"][f] > 0, "4+ count %s positive" % f)
    T = g["T"]
    cur = g["cur"]
    ok(len(g["meds"]) == 2 and all(cur[d] >= 1.5 * T for d in g["meds"]), "two major event days, each at 1.5 times "
                                                                           "the threshold or more")
    ok(max(v for d, v in cur.items() if d not in g["meds"]) <= 0.75 * T, "every other day at 0.75 of the threshold "
                                                                         "or less")
    T0 = A.t_med(g["prior"], ddof=0)
    ok(sorted(d for d, v in cur.items() if v > T0) == g["meds"], "population and sample standard deviation give the "
                                                                  "same major event days")
    ok([d.isoformat() for d in g["meds"]] == ["2026-01-25", "2026-07-28"], "the major event days are the ice storm and "
                                                                           "the July thunderstorm")
    # every substation has a non-MED bank outage in the asked year, and the 30 September evening sits past midnight UTC
    banks = oms[(oms["device_type"] == "TRANSFORMER BANK")]
    bank_days = set()
    for rr in banks.itertuples(index=False):
        d = A.parse(rr.start_ts, rr.source_system).date()
        if A.RY[0] <= d <= A.RY[1] and d not in g["meds"]:
            bank_days.add(rr.feeder_id)
    ok(bank_days == set(W.FEEDERS), "every feeder had a substation bank outage on an ordinary day in the asked year")
    late = oms[(oms["source_system"] == "ADMS") & oms["start_ts"].str.startswith("2026-10-01")]
    ok(len(late) >= 3, "outages of the evening of 30 September 2026 carry 1 October UTC stamps")
    # the outage records' own integrity
    ok(not oms.duplicated(["outage_id", "feeder_id"]).any(), "one export row per outage per feeder")
    ok(set(steps["outage_id"]) == set(oms["outage_id"]), "every outage has its steps")
    sm = steps.groupby("outage_id")["customers_restored"].sum()
    co = oms.groupby("outage_id")["customers_out"].first()
    ok((sm == co.reindex(sm.index)).all(), "steps' customers restored sum to the outage's customers out")
    intr = g["intr"]
    ok(not ((intr["own_min"] > 4.0) & (intr["own_min"] <= 6.0)).any(), "no customer restoration within a minute of "
                                                                       "the five-minute line")
    legacy = oms[oms["source_system"] == "OMS"]
    ok(legacy["start_ts"].max() < "2026-03-17 06:00:00" and
       oms[oms["source_system"] == "ADMS"]["start_ts"].min() > "2026-03-17 10:00:00", "the systems split at the cutover")
    # the designed stops: the cracker misses every figure; every stop lands off the golden on the feeders it moves
    S = A.stops(g, oms, steps, conn, dv, cis)
    ck = S["A1 cracker: export columns summed, connectivity customers"]
    for f in W.FEEDERS:
        ok(round(ck[0][f], 2) != round(g["saifi"][f], 2), "A1 cracker misses SAIFI %s" % f)
        ok(round(ck[1][f], 1) != round(g["saidi"][f], 1), "A1 cracker misses SAIDI %s" % f)
        ok(S["A2 cracker: distinct outages, outage duration"][f] != g["cemi4"][f], "A2 cracker misses %s" % f)
        nat = S["A1 natural: export columns summed, customer-system labels"]
        ok(round(nat[0][f], 2) != round(g["saifi"][f], 2) and round(nat[1][f], 1) != round(g["saidi"][f], 1),
           "A1 natural stop misses %s" % f)
        ok(S["A2 rows: one interruption per step row"][f] != g["cemi4"][f] or f == "SP09", "A2 rows stop misses %s" % f)
    lab = S["A1 labels: each customer's restoration, customer-system labels"]
    ok(all(round(lab[0][f], 2) != round(g["saifi"][f], 2) for f in ("ML03", "ML07")), "the stale labels move the "
                                                                                       "Mill Lane pair")
    rf = S["A2 row feeder: each customer's restoration, the outage's feeder"]
    ok(rf["ML03"] != g["cemi4"]["ML03"] or rf["ML07"] != g["cemi4"]["ML07"], "the outage's feeder splits the moved "
                                                                              "lateral's customers")
    clk = S["A1 clock: each customer's restoration, ADMS times read as local"]
    moved_clk = [f for f in W.FEEDERS if round(clk[1][f], 1) != round(g["saidi"][f], 1)]
    ok(len(moved_clk) >= 2, "reading ADMS times as local moves SAIDI on two feeders or more: %s" % moved_clk)
    for k, v in S.items():
        if isinstance(v, tuple):
            LOG.append("%-70s SAIFI %s | SAIDI %s" % (k, " ".join("%.2f" % v[0][f] for f in W.FEEDERS),
                                                      " ".join("%.1f" % v[1][f] for f in W.FEEDERS)))
        else:
            LOG.append("%-70s %s" % (k, " ".join(str(v[f]) for f in W.FEEDERS)))
    LOG.append("%-70s SAIFI %s | SAIDI %s | 4+ %s" % ("golden", " ".join("%.2f" % g["saifi"][f] for f in W.FEEDERS),
                                                       " ".join("%.1f" % g["saidi"][f] for f in W.FEEDERS),
                                                       " ".join(str(g["cemi4"][f]) for f in W.FEEDERS)))
    # the published major event days reproduce on the shipped records
    pub = {}
    for rr in meds_pub.itertuples(index=False):
        pub.setdefault(rr.reliability_year, (float(rr.tmed_saidi_min), set()))[1].add(dt.date.fromisoformat(rr.med_date))
    for y in range(2021, 2026):
        Ty, days = pub["RY%d" % y]
        dl = A.daily_saidi(intr, g["net"], [(dt.date(y - 1, 10, 1), dt.date(y, 9, 30))])
        ok({d for d, v in dl.items() if v > Ty} == days, "published RY%d major event days reproduce" % y)

    # ------------------------------------------------------------ separation: ask devices off the main call's rows
    main_files = {R.WO, R.REGISTER, R.CODES}
    device_files = {R.OMS, R.STEPS, R.CONN, R.CIS, R.DEVICES, R.MED}
    ok(not (main_files & device_files), "zero device rows inside the main call's population (disjoint files)")
    ok(not set(wo.columns) & {"outage_id", "meter_id", "customers_out", "start_ts"}, "the work orders carry no outage "
                                                                                     "or meter field")

    # ------------------------------------------------------------ distractors unused
    meta = json.load(open(meta_path)) if meta_path else None
    if meta:
        ok(set(meta["distractor_files"]) == {R.VAULTS, R.LOADING}, "metadata names the two distractors")
        tmp = tempfile.mkdtemp(dir=TASK, suffix=".tmp")
        try:
            t2 = os.path.join(tmp, "t")
            shutil.copytree(target, t2)
            for f in meta["distractor_files"]:
                os.remove(os.path.join(t2, f))
            c2 = DL.compute(t2)
            ok(c2["order"][0] == "QH21" and json.dumps(DL.figures(c2), sort_keys=True) ==
               json.dumps(figs, sort_keys=True), "the golden recomputes unchanged with the distractors removed")
        finally:
            shutil.rmtree(tmp)
        ok(meta["input_gates"]["files"] >= 10 and len(meta["input_gates"]["formats"]) >= 3, "input gates: files, formats")
        ok(meta["input_gates"]["largest_file_rows"] >= 25000, "input gates: a file of 25,000 rows or more")

    # ------------------------------------------------------------ pack gates, names, vocabulary, dates
    files = sorted(os.listdir(target))
    ok(len(files) == 14 and len({f.rsplit(".", 1)[1] for f in files}) >= 5, "14 files in 5 or more formats")
    ok(not any("distractor" in f.lower() for f in files), "no file named for a distractor")
    for f, t in texts:
        ok("distractor" not in t.lower() and "—" not in t, "no distractor word or em dash in %s" % f)
    ok(len(conn) >= 25000, "connectivity model has 25,000 rows or more: %d" % len(conn))
    gx = openpyxl.load_workbook(os.path.join(golden, DL.XLSX))
    ws = gx["Feeders"]
    rows = {ws.cell(row=i, column=1).value: [ws.cell(row=i, column=j).value for j in range(1, 11)] for i in range(2, 10)}
    for f in W.FEEDERS:
        rr = rows[W.NAME[f]]
        ok("%.2f" % rr[7] == figs["saifi_ry2026"][f] and "%.1f" % rr[8] == figs["saidi_ry2026_min"][f] and
           rr[9] == figs["customers_4plus_ry2026"][f], "the workbook carries %s's figures" % f)
    ok(gx["2027 program"]["B4"].value == "Quarry Hill 21", "the workbook names Quarry Hill 21")
    ok(os.path.getsize(os.path.join(golden, DL.PNG)) > 20000, "the chart is rendered")
    for p in (os.path.join(golden, DL.XLSX),) + tuple(os.path.join(target, f) for f in files if f.endswith(".xlsx")):
        with zipfile.ZipFile(p) as z:
            core = z.read("docProps/core.xml").decode()
            ok("openpyxl" not in core and "openpyxl" not in z.read("docProps/app.xml").decode(), "no writer library in %s" % p)
    for f in files:
        if f.endswith(".pdf"):
            b = open(os.path.join(target, f), "rb").read()
            ok(b"ReportLab" not in b and b"reportlab" not in b, "no writer library in %s" % f)
    people = set(W.PEOPLE.values())
    card = json.load(open(os.path.join(REPO, ".claude", "skills", "fingerprint", "cards", "task183.json")))
    ok(people == set(card["world"]["people"]), "the pack's people are the card's")
    notes_t = open(os.path.join(target, "asset_analytics_extract_notes_2026-10.md")).read()
    ok(all(p.split()[1] in notes_t + "".join(t for _, t in texts) for p in people), "the people appear in the pack")
    # the hand-off note keeps the conventions the files need read literally (replacement, clocks, dated rows)
    for must in ("new section with predecessor_section_id", "completed_date is the day the new\n  cable was energised",
                 "Legacy OMS\n  times are Eastern local, ADMS times are UTC", "eff_to is the first day the row no longer applies"):
        ok(must in notes_t, "the file notes keep: %s" % must.replace("\n  ", " "))
    ok(max(conn["eff_from"]) <= W.CUT, "no connectivity row starts after the cut")
    # the prompt and the voice
    prompt = open(os.path.join(TASK, "prompt.txt")).read()
    ok(len(prompt) < 1000 and "—" not in prompt, "prompt under 1,000 characters, no em dash")

    # ------------------------------------------------------------ chart order stability
    orders = {name: lead(fc)[3] for name, fc in variants.items()}
    tops = {o[0] for o in orders.values()}
    full = {tuple(o) for o in orders.values()}
    ok(tops == {"QH21"}, "every R3 variant puts Quarry Hill 21 on top of the chart")
    LOG.append("chart order: %d distinct full orders across %d R3 variants (top feeder stable, full order not)" % (
        len(full), len(orders)))
    if not quiet:
        for line in LOG:
            print(line)
    print("checks: %d assertions passed" % N[0])
    return N[0]


def report_check(target, wo, reg, conn):
    with pdfplumber.open(os.path.join(target, "2025_annual_reliability_report_sec6_ug_cable.pdf")) as pdf:
        text = "\n".join(p.extract_text() for p in pdf.pages)
    years = [2021, 2022, 2023, 2024, 2025]
    good = True
    for f in W.FEEDERS:
        w = wo[(wo["feeder_id"] == f) & wo["cause_code"].isin(M.CABLE)]
        counts = [int((w["fault_date"].map(lambda d: d.year) == y).sum()) for y in years]
        line = "%s %s %d" % (W.NAME[f], " ".join(str(c) for c in counts), sum(counts))
        good &= line in text
        cust = []
        for y in years:
            d = dt.date(y, 12, 31)
            live = conn[(conn["eff_from"] <= d) & conn["eff_to"].map(lambda x: x is None or d < x)]
            cust.append(int(live[live["feeder_id"] == f]["meter_id"].nunique()))
        good &= ("%s %s" % (W.NAME[f], " ".join("{:,}".format(c) for c in cust))) in text
    return good


if __name__ == "__main__":
    run(os.path.join(TASK, "target"), os.path.join(TASK, "golden"), meta_path=os.path.join(TASK, "metadata.json"))
