"""The goldens, computed from the pack in target/: cable_program_2027.xlsx (the named feeder, every feeder's 2027
expected cable failures and its RY2026 reliability figures) and feeder_failures_2027.png (the bar chart), plus
generator/figures.json, the owed figures the write-up is drafted from.

The 2027 forecast: each in-service section's failure rate for its count of cable failures so far (0, 1, 2, 3 or
more), estimated on every section year of the record (January 2012 to 30 September 2026, exact time at risk),
summed by feeder."""
import datetime as dt
import json
import os
import sys

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt   # noqa: E402
import openpyxl                   # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side   # noqa: E402
from openpyxl.utils import get_column_letter                             # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import asks as A          # noqa: E402
import load as L          # noqa: E402
import models as M        # noqa: E402
import postprocess as PP  # noqa: E402
import world as W         # noqa: E402

XLSX = "cable_program_2027.xlsx"
PNG = "feeder_failures_2027.png"
INK, INK2, MUTED, ACCENT, SURFACE, GRID = "#0b0b0b", "#52514e", "#b9b7b0", "#2a78d6", "#fcfcfb", "#e4e2dc"


def compute(target):
    wo, reg = L.work_orders(target), L.register(target)
    fc, rates = M.kmodel(wo, reg)
    _, fails = M.prep(wo, reg)
    ex = M.exposures(reg, fails)
    ex["kc"] = ex["k"].clip(upper=3)
    tab = ex.groupby("kc")[["years", "events"]].sum()
    st = M.stock(reg, fails)
    twice = st[st["k"] >= 2].groupby("feeder_id").size().to_dict()
    n_sec = st.groupby("feeder_id").size().to_dict()
    oms, steps, conn, dv = L.oms(target), L.steps(target), L.connectivity(target), L.devices(target)
    g = A.golden(oms, steps, conn, dv)
    order, lead = M.rank(fc)
    return dict(fc=fc, rates=rates, tab=tab, twice=twice, n_sec=n_sec, g=g, order=order, lead=lead)


def figures(c):
    g = c["g"]
    out = {
        "recommended_feeder": W.NAME[c["order"][0]],
        "recommended_feeder_id": c["order"][0],
        "sections_twice_failed_in_service": {f: int(c["twice"].get(f, 0)) for f in W.FEEDERS},
        "saifi_ry2026": {f: "%.2f" % g["saifi"][f] for f in W.FEEDERS},
        "saidi_ry2026_min": {f: "%.1f" % g["saidi"][f] for f in W.FEEDERS},
        "customers_4plus_ry2026": {f: int(g["cemi4"][f]) for f in W.FEEDERS},
        "customers_served_2026_09_30": {f: int(g["served"][f]) for f in W.FEEDERS},
        "major_event_days_ry2026": [d.isoformat() for d in g["meds"]],
        "major_event_day_method": "2.5 beta (TVE-RS-04)",
        "feeder_names": [W.NAME[f] for f in W.FEEDERS],
        "twice_failed_sections_location": "State Route 443",
    }
    return out


def write_xlsx(path, c):
    g = c["g"]
    wb = openpyxl.Workbook()
    thin = Side(style="thin", color="A6A39B")
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    pick_fill = PatternFill("solid", fgColor="E3EEFA")

    ws = wb.active
    ws.title = "2027 program"
    ws["A1"] = "2027 underground cable program: Quarry Hill 21"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Distribution Asset Management, Tamsin Valley district. Prepared %s from the October 2026 extract " \
               "(records through 30 September 2026)." % "9 October 2026"
    ws["A2"].font = Font(italic=True, size=9, color="52514E")
    lines = [
        ("Feeder for the 2027 program", W.NAME[c["order"][0]]),
        ("Basis", "Expected failures of the cable itself (insulation, splice, termination) in calendar 2027 on the "
                  "sections in service at 30 September 2026, damage by others left out"),
        ("Method", "Each section's failure rate for the number of cable failures it has already had (0, 1, 2, 3 or "
                   "more), estimated on every section year since January 2012, summed over the feeder's sections"),
        ("Why not the five-year counts", "Mill Lane 7's recent failures are first failures on different sections, "
                                         "which rarely fail again; Quarry Hill 21 is the only feeder with in-service "
                                         "sections that have failed twice or more (the State Route 443 crossing)"),
    ]
    for i, (k, v) in enumerate(lines, 4):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        cell = ws.cell(row=i, column=2, value=v)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=1).alignment = Alignment(vertical="top")
    ws.cell(row=4, column=2).font = Font(bold=True, size=12, color="1F3A5F")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 96
    for r in (5, 6, 7):
        ws.row_dimensions[r].height = 32

    ws = wb.create_sheet("Feeders")
    cols = ["Feeder", "Circuit", "Substation", "UG sections in service", "In-service sections with 2+ failures",
            "Expected cable failures 2027", "Customers served 30 Sep 2026", "SAIFI RY2026 (MED excl.)",
            "SAIDI RY2026, min (MED excl.)", "Customers with 4+ sustained interruptions"]
    for j, h in enumerate(cols, 1):
        cell = ws.cell(row=1, column=j, value=h)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 44
    sub = {f: n for s, n, fs in W.SUBSTATIONS for f in fs}
    for i, f in enumerate(W.FEEDERS, 2):
        vals = [W.NAME[f], W.CIRCUIT_NO[f], sub[f], int(c["n_sec"].get(f, 0)), int(c["twice"].get(f, 0)),
                round(c["fc"][f], 1), int(g["served"][f]), round(g["saifi"][f], 2), round(g["saidi"][f], 1),
                int(g["cemi4"][f])]
        fmts = [None, None, None, "#,##0", "0", "0.0", "#,##0", "0.00", "#,##0.0", "#,##0"]
        for j, (v, fm) in enumerate(zip(vals, fmts), 1):
            cell = ws.cell(row=i, column=j, value=v)
            if fm:
                cell.number_format = fm
                cell.alignment = Alignment(horizontal="right")
            cell.border = Border(bottom=Side(style="hair", color="D0CEC7"))
            if f == c["order"][0]:
                cell.fill = pick_fill
                cell.font = Font(bold=True)
    tot = len(W.FEEDERS) + 2
    ws.cell(row=tot, column=1, value="District").font = Font(bold=True)
    for j, v, fm in ((4, sum(c["n_sec"].values()), "#,##0"), (5, sum(c["twice"].values()), "0"),
                     (7, sum(g["served"].values()), "#,##0")):
        cell = ws.cell(row=tot, column=j, value=int(v))
        cell.number_format = fm
        cell.font = Font(bold=True)
        cell.border = Border(top=thin)
    ws.cell(row=tot + 2, column=1, value="Expected failures are model estimates for ranking the feeders; the SAIFI, "
                                          "SAIDI and 4+ counts are the reliability year 1 Oct 2025 to 30 Sep 2026 "
                                          "under TVE-RS-04 rev 6.").font = Font(italic=True, size=9, color="52514E")
    widths = [16, 9, 12, 11, 13, 12, 13, 11, 12, 14]
    for j, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = wdt
    ws.freeze_panes = "B2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    ws = wb.create_sheet("Failure rates")
    for j, h in enumerate(["Cable failures already on the section", "Section years at risk", "Failures",
                           "Failures per section year"], 1):
        cell = ws.cell(row=1, column=j, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    for i, (kc, row) in enumerate(c["tab"].iterrows(), 2):
        ws.cell(row=i, column=1, value=["None", "One", "Two", "Three or more"][int(kc)])
        ws.cell(row=i, column=2, value=round(float(row["years"]), 1)).number_format = "#,##0.0"
        ws.cell(row=i, column=3, value=int(row["events"])).number_format = "#,##0"
        ws.cell(row=i, column=4, value=round(float(row["events"] / row["years"]), 4)).number_format = "0.0%"
    ws.cell(row=7, column=1, value="Section years from January 2012 (start of the work order archive) to 30 September "
                                   "2026, split at each failure; INS, SPL and TRM work orders only.").font = Font(
        italic=True, size=9, color="52514E")
    for j, wdt in enumerate([30, 14, 10, 14], 1):
        ws.column_dimensions[get_column_letter(j)].width = wdt

    ws = wb.create_sheet("Notes")
    notes = [
        "Cable failures: work orders with cause INS, SPL or TRM (wo_cause_repair_codes.csv); TPD (damage by others) "
        "and OTH left out.",
        "Sections in service: status IS in the GIS export as of 30 September 2026; a section's failures stay with "
        "it, and a replacement section starts its own record.",
        "Reliability year RY2026: 1 October 2025 to 30 September 2026, Eastern local time. ADMS times (UTC) "
        "converted before assigning days.",
        "Sustained interruption: more than five minutes for the customer, timed to that customer's own restoration "
        "step (TVE-RS-04 3.3); customers served from the connectivity model on 30 September 2026 (3.2).",
        "Major event days RY2026 (2.5 beta, threshold %.2f min of daily SAIDI from RY2021 to RY2025): %s. Left out "
        "of SAIFI and SAIDI, kept in the 4+ count (TVE-RS-04 6.1)." % (g["T"], ", ".join(
            d.strftime("%d %b %Y").lstrip("0") for d in g["meds"])),
        "Substation bank outages counted once per customer on each feeder, not at the outage totals the OMS export "
        "repeats on each feeder row.",
    ]
    ws["A1"] = "Notes"
    ws["A1"].font = Font(bold=True)
    for i, t in enumerate(notes, 2):
        ws.cell(row=i, column=1, value="%d. %s" % (i - 1, t)).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 30
    ws.column_dimensions["A"].width = 120
    wb.active = 0
    wb.save(path)
    PP.normalize_ooxml(path, W.PEOPLE["asset"], W.PEOPLE["asset"], "2026-10-07T13:12:40Z", "2026-10-09T20:47:15Z",
                       excel=True)


def write_png(path, c):
    fc = c["fc"]
    order = sorted(W.FEEDERS, key=lambda f: fc[f])
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9.5, "svg.hashsalt": "t183",
                         "axes.edgecolor": INK2, "axes.labelcolor": INK2, "xtick.color": INK2, "ytick.color": INK})
    fig, ax = plt.subplots(figsize=(7.6, 4.3), dpi=150)
    fig.patch.set_facecolor(SURFACE)
    ax.set_facecolor(SURFACE)
    colors = [ACCENT if f == c["order"][0] else MUTED for f in order]
    ax.barh([W.NAME[f] for f in order], [fc[f] for f in order], color=colors, height=0.62, edgecolor=SURFACE,
            linewidth=2)
    for i, f in enumerate(order):
        ax.text(fc[f] + 0.15, i, "%.1f" % fc[f], va="center", fontsize=8.5,
                color=INK if f == c["order"][0] else INK2, fontweight="bold" if f == c["order"][0] else "normal")
    top = c["order"][0]
    n2 = int(c["twice"].get(top, 0))
    second = sorted(fc.values())[-2]
    ax.annotate("%d sections already failed\ntwice or more, spliced in place\nunder State Route 443" % n2,
                xy=(second + 2.6, len(order) - 1.3), xytext=(second + 1.4, len(order) - 2.55),
                fontsize=7.8, color=INK2, ha="left", va="center", linespacing=1.3,
                arrowprops=dict(arrowstyle="-", color=INK2, linewidth=0.7, shrinkA=2, shrinkB=0))
    ax.set_xlabel("Expected cable failures in 2027 (insulation, splice, termination)")
    ax.xaxis.grid(True, color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.tick_params(axis="y", length=0)
    ax.set_xlim(0, max(fc.values()) * 1.12)
    fig.suptitle("Quarry Hill 21 is expected to have the most cable failures in 2027", x=0.02, ha="left",
                 fontsize=11.5, fontweight="bold", color=INK)
    ax.set_title("Sections in service at 30 Sep 2026, each at the failure rate for the failures it has already had",
                 loc="left", fontsize=8.5, color=INK2)
    fig.text(0.02, 0.015, "Source: cable fault work orders 2012-2026 and GIS UG primary sections, Tamsin Valley district",
             fontsize=7, color=INK2)
    fig.tight_layout(rect=(0, 0.03, 1, 0.97))
    fig.savefig(path, facecolor=SURFACE, metadata={"Software": None})
    plt.close(fig)


def main(target, golden, figures_path):
    os.makedirs(golden, exist_ok=True)
    c = compute(target)
    write_xlsx(os.path.join(golden, XLSX), c)
    write_png(os.path.join(golden, PNG), c)
    with open(figures_path, "w") as fh:
        json.dump(figures(c), fh, indent=1)
        fh.write("\n")
    return c


if __name__ == "__main__":
    task = os.path.dirname(HERE)
    main(os.path.join(task, "target"), os.path.join(task, "golden"), os.path.join(HERE, "figures.json"))
