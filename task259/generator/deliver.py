"""The goldens: the band workbook the head of total rewards files and the chart beside it, and figures.json, every
figure the write-up states, computed once from the pack as shipped."""
from __future__ import annotations

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties

import analysis as A
import common as K
import load as L
import postprocess as PP

ZONES = (1, 2, 3)


def compute(target):
    ex, pr, cw = L.extract(target), L.profiles(target), L.crosswalk(target)
    fam, focus = L.survey_match(target)
    G = A.golden(ex, pr, cw, fam, focus)
    R = A.rungs(ex, pr, cw, fam, focus)
    hr, zn, req = L.hris(target), L.zones(target), L.ats(target)
    hc = A.headcount(hr, zn)
    oreq = A.open_requisitions(req)
    f = {}
    f["call_DE_III_midpoint"] = G["filed"]["III"][50]
    for lv in A.LEVELS:
        f["midpoint_%s" % lv] = G["filed"][lv][50]
        f["matched_%s" % lv] = G["n"][lv]
        f["open_requisitions_%s" % lv] = oreq[lv]
        for z in ZONES:
            f["headcount_%s_zone%d" % (lv, z)] = int(hc.loc[lv, z])
    f["DE_III_p25"] = G["filed"]["III"][25]
    f["DE_III_p75"] = G["filed"]["III"][75]
    f["movement_pct_a_year"] = round(100 * G["g"], 1)
    f["pooled_movement_pct_a_year"] = round(100 * R["g_pooled"], 1)
    f["paired_growth_pct_a_year"] = round(100 * R["g_pair"], 1)
    f["DE_III_unaged"] = A.to_k(R["R1"]["III"][50])
    f["DE_III_pooled_movement"] = A.to_k(R["R3"]["III"][50])
    f["DE_III_paired_growth"] = A.to_k(R["R2"]["III"][50])
    f["matched_total"] = int(sum(G["n"].values()))
    f["headcount_total"] = int(hc.to_numpy().sum())
    f["open_requisitions_total"] = int(sum(oreq.values()))
    return dict(figures=f, G=G, R=R, headcount=hc, open_reqs=oreq)


THIN = Side(style="thin", color="B9B4C2")
HEAD_FILL = PatternFill("solid", fgColor="3E2A4F")
CALL_FILL = PatternFill("solid", fgColor="EFE8F5")
MONEY = '"$"#,##0'


def _head(ws, row, values, widths=None):
    for j, v in enumerate(values, 1):
        c = ws.cell(row=row, column=j, value=v)
        c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center" if j > 1 else "left", vertical="center", wrap_text=True)
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + j)].width = w


def write_workbook(path, fig):
    f, G = fig["figures"], fig["G"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Bands 2027"
    ws["A1"] = "Data Engineer pay bands from 1 April 2027"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True)
    ws["A2"] = ("Market as in effect on 1 October 2026: matched Tallyband respondents, each report carried to the "
                "benchmark date at the movement of pay at a fixed level within employers (%.1f per cent a year). "
                "Filed to the nearest $1,000." % f["movement_pct_a_year"])
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="5D5A66")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:E2")
    ws.row_dimensions[2].height = 30
    _head(ws, 4, ["Level", "25th percentile", "Midpoint", "75th percentile", "Matched respondents"],
          (24, 17, 17, 17, 21))
    for i, lv in enumerate(A.LEVELS):
        r = 5 + i
        ws.cell(row=r, column=1, value="Data Engineer %s" % lv)
        ws.cell(row=r, column=3, value=f["midpoint_%s" % lv]).number_format = MONEY
        if lv == "III":
            ws.cell(row=r, column=2, value=f["DE_III_p25"]).number_format = MONEY
            ws.cell(row=r, column=4, value=f["DE_III_p75"]).number_format = MONEY
            for j in range(1, 6):
                ws.cell(row=r, column=j).fill = CALL_FILL
                ws.cell(row=r, column=j).font = Font(name="Calibri", size=11, bold=True)
        ws.cell(row=r, column=5, value=f["matched_%s" % lv]).number_format = "#,##0"
        for j in range(1, 6):
            ws.cell(row=r, column=j).border = Border(bottom=THIN)
            if j > 1:
                ws.cell(row=r, column=j).alignment = Alignment(horizontal="right")
    ws.cell(row=9, column=1, value="Total").font = Font(name="Calibri", size=11, bold=True)
    ws.cell(row=9, column=5, value=f["matched_total"]).number_format = "#,##0"
    ws.cell(row=9, column=5).font = Font(name="Calibri", size=11, bold=True)
    ws["A11"] = ("The other levels' 25th and 75th percentiles are drawn on data_engineer_bands_2027.svg. Policy 4.4 "
                 "files each point to the nearest $1,000; the points are read from unrounded pay.")
    ws["A11"].font = Font(name="Calibri", size=9, color="5D5A66")
    ws["A11"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A11:E11")
    ws.row_dimensions[11].height = 28
    ws.freeze_panes = "B5"

    hc = wb.create_sheet("Headcount 2026-10-01")
    hc["A1"] = "Data Engineers employed on 1 October 2026, by level and pay zone"
    hc["A1"].font = Font(name="Calibri", size=13, bold=True)
    _head(hc, 3, ["Level", "Zone 1", "Zone 2", "Zone 3", "All zones"], (24, 12, 12, 12, 13))
    for i, lv in enumerate(A.LEVELS):
        r = 4 + i
        hc.cell(row=r, column=1, value="Data Engineer %s" % lv)
        for j, z in enumerate(ZONES, 2):
            hc.cell(row=r, column=j, value=f["headcount_%s_zone%d" % (lv, z)])
        hc.cell(row=r, column=5, value=sum(f["headcount_%s_zone%d" % (lv, z)] for z in ZONES)).font = Font(bold=True)
    hc.cell(row=8, column=1, value="All levels").font = Font(name="Calibri", size=11, bold=True)
    for j, z in enumerate(ZONES, 2):
        hc.cell(row=8, column=j, value=sum(f["headcount_%s_zone%d" % (lv, z)] for lv in A.LEVELS)).font = Font(bold=True)
    hc.cell(row=8, column=5, value=f["headcount_total"]).font = Font(bold=True)
    for r in range(4, 9):
        for j in range(1, 6):
            hc.cell(row=r, column=j).border = Border(bottom=THIN)
    hc["A10"] = ("Employees only (contingent workers excluded), on leave included; each employee's job record in "
                 "effect on 1 October 2026; zone of the work location under the zone map in force that day.")
    hc["A10"].font = Font(name="Calibri", size=9, color="5D5A66")
    hc["A10"].alignment = Alignment(wrap_text=True, vertical="top")
    hc.merge_cells("A10:E10")
    hc.row_dimensions[10].height = 40

    rq = wb.create_sheet("Open requisitions")
    rq["A1"] = "Data Engineer requisitions open on 1 October 2026"
    rq["A1"].font = Font(name="Calibri", size=13, bold=True)
    _head(rq, 3, ["Level", "Open requisitions"], (24, 19))
    for i, lv in enumerate(A.LEVELS):
        rq.cell(row=4 + i, column=1, value="Data Engineer %s" % lv)
        rq.cell(row=4 + i, column=2, value=f["open_requisitions_%s" % lv])
        for j in (1, 2):
            rq.cell(row=4 + i, column=j).border = Border(bottom=THIN)
    rq.cell(row=8, column=1, value="All levels").font = Font(bold=True)
    rq.cell(row=8, column=2, value=f["open_requisitions_total"]).font = Font(bold=True)
    rq["A10"] = "Each requisition once, whatever the number of cities it is posted in."
    rq["A10"].font = Font(name="Calibri", size=9, color="5D5A66")

    nt = wb.create_sheet("Notes")
    for line in ("Prepared by %s, compensation analytics, for %s, head of total rewards." % (K.PEOPLE["comp"],
                                                                                            K.PEOPLE["head"]),
                 "Benchmark date 1 October 2026 (policy 4.3); bands take effect 1 April 2027 (policy 4.5).",
                 "Sources: Tallyband extract and profile table of 1 October 2026; level crosswalk 2026; engineering "
                 "job architecture; HRIS job history and ATS requisitions exported 14 October 2026; pay zones 2026."):
        nt.append([line])
    nt.column_dimensions["A"].width = 110
    for s in wb.worksheets:
        s.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        s.page_setup.orientation = "landscape"
    wb.save(path)
    PP.normalize_ooxml(path, K.PEOPLE["comp"], K.PEOPLE["comp"], "2026-10-15T21:48:03Z", "2026-10-16T17:12:40Z",
                       excel=True)


def _money(v):
    return "${:,.0f}".format(v)


def write_chart(path, fig):
    """Each level's band from its 25th to its 75th percentile with its midpoint marked; Data Engineer III labelled
    in full. Hand-drawn SVG, so the file carries no tool stamp."""
    f, P = fig["figures"], fig["G"]["points"]
    W, H, left, right, top = 900, 430, 170, 40, 92
    lo, hi = 100_000, 230_000
    x = lambda v: left + (v - lo) / (hi - lo) * (W - left - right)
    rowh = 62
    out = ['<svg xmlns="http://www.w3.org/2000/svg" width="%d" height="%d" viewBox="0 0 %d %d" '
           'font-family="Helvetica, Arial, sans-serif">' % (W, H, W, H),
           '<rect width="%d" height="%d" fill="#ffffff"/>' % (W, H),
           '<text x="24" y="38" font-size="20" font-weight="bold" fill="#1f1d24">Data Engineer III midpoint '
           '%s from 1 April 2027</text>' % _money(f["call_DE_III_midpoint"]),
           '<text x="24" y="62" font-size="13" fill="#5d5a66">Bands from the market as in effect on 1 October 2026: '
           '25th to 75th percentile, midpoint marked</text>']
    for t in range(lo, hi + 1, 10_000):
        xx = x(t)
        out.append('<line x1="%.1f" y1="%d" x2="%.1f" y2="%d" stroke="#e6e2ec" stroke-width="1"/>' % (
            xx, top - 6, xx, top + rowh * 4))
        if t % 20_000 == 0:
            out.append('<text x="%.1f" y="%d" font-size="11" fill="#5d5a66" text-anchor="middle">$%dk</text>' % (
                xx, top + rowh * 4 + 18, t // 1000))
    for i, lv in enumerate(A.LEVELS):
        y = top + i * rowh + 14
        call = lv == "III"
        fill = "#6b3fa0" if call else "#b9a7cf"
        a, m, b = P[lv][25], P[lv][50], P[lv][75]
        out.append('<text x="24" y="%d" font-size="14" fill="#1f1d24"%s>Data Engineer %s</text>' % (
            y + 19, ' font-weight="bold"' if call else "", lv))
        out.append('<rect x="%.1f" y="%d" width="%.1f" height="28" rx="3" fill="%s"/>' % (x(a), y, x(b) - x(a), fill))
        out.append('<line x1="%.1f" y1="%d" x2="%.1f" y2="%d" stroke="#1f1d24" stroke-width="3"/>' % (
            x(m), y - 5, x(m), y + 33))
        out.append('<text x="%.1f" y="%d" font-size="12" fill="#1f1d24" text-anchor="middle"%s>%s</text>' % (
            x(m), y - 9, ' font-weight="bold"' if call else "", _money(f["midpoint_%s" % lv])))
        if call:
            out.append('<text x="%.1f" y="%d" font-size="11" fill="#1f1d24" text-anchor="end">%s</text>' % (
                x(a) - 6, y + 19, _money(f["DE_III_p25"])))
            out.append('<text x="%.1f" y="%d" font-size="11" fill="#1f1d24">%s</text>' % (
                x(b) + 6, y + 19, _money(f["DE_III_p75"])))
    out.append('<text x="24" y="%d" font-size="10.5" fill="#5d5a66">Source: Tallyband extract of 1 October 2026, '
               'matched respondents (policy 4.2), each report carried to the benchmark date at %.1f per cent a year.'
               '</text>' % (H - 18, f["movement_pct_a_year"]))
    out.append("</svg>")
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write("\n".join(out) + "\n")
