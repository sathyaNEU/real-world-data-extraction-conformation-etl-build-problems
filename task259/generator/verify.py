"""Independent verifier: reads only the shipped pack and the goldens, recomputes every graded figure by its own code
path (csv rows, its own name tidy, its own crosswalk lookup, statsmodels for the slope, its own as-of reads) and
compares with the band workbook and the chart.

    python3 task259/generator/verify.py <target> <golden>"""
from __future__ import annotations

import csv
import datetime as dt
import os
import re
import sys

import numpy as np
import openpyxl
import statsmodels.api as sm

BENCH = dt.date(2026, 10, 1)
LV = ("I", "II", "III", "IV")


def _rows(target, name):
    with open(os.path.join(target, name), encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def _d(s):
    return dt.date.fromisoformat(s) if s else None


def _key(name):
    s = re.sub(r"[.,]", " ", name.lower()).strip()
    s = re.sub(r"\binc$", "", s).strip()
    return re.sub(r"[^a-z0-9]+", "", s)


def bands(target):
    wb = openpyxl.load_workbook(os.path.join(target, "tarnwick_engineering_job_architecture.xlsx"), data_only=True)
    jobs = list(wb["Jobs"].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(jobs) if r and r[0] == "Family")
    head = list(jobs[hi])
    de = next(dict(zip(head, r)) for r in jobs[hi + 1:] if r[1] == "Data Engineer")
    fam, focus = de["Tallyband family"], de["Tallyband primary focus"]
    focus_of = {r["respondent_id"]: r["primary_focus"] for r in _rows(target, "tallyband_profiles_2026-10-01.csv")}
    xw = {}
    for r in _rows(target, "level_crosswalk_2026.csv"):
        xw.setdefault((_key(r["employer"]), r["level_code"]), []).append(
            (_d(r["effective_from"]), _d(r["effective_to"]), r["tarnwick_level"]))
    recs = []
    for r in _rows(target, "tallyband_extract_2026-10-01.csv"):
        if r["country"] != "US" or r["employment_type"] != "full_time" or r["job_family"] != fam:
            continue
        if focus_of.get(r["respondent_id"]) != focus:
            continue
        filed = _d(r["filed_on"])
        lv = [t for a, b, t in xw.get((_key(r["employer"]), r["employer_level"]), ())
              if a <= filed and (b is None or filed <= b)]
        if lv and lv[0] in LV:
            recs.append((lv[0], filed, float(r["base_salary_usd"]), _key(r["employer"])))
    lev = np.array([x[0] for x in recs])
    yrs = np.array([(BENCH - x[1]).days / 365.25 for x in recs])
    pay = np.array([x[2] for x in recs])
    emp = np.array([x[3] for x in recs])
    names = sorted(set(emp))
    X = np.column_stack([(lev == l).astype(float) for l in LV] + [(emp == e).astype(float) for e in names[1:]] +
                        [-yrs])
    g = sm.OLS(np.log(pay), X).fit().params[-1]           # within employers: an intercept for each employer
    carried = pay * np.exp(g * yrs)
    out = {}
    for l in LV:
        v = carried[lev == l]
        out[l] = {q: int(np.floor(np.percentile(v, q) / 1000 + 0.5) * 1000) for q in (25, 50, 75)}
    return out, g, {l: int((lev == l).sum()) for l in LV}


def headcount(target):
    zones = {}
    for r in _rows(target, "pay_zones_2026.csv"):
        zones.setdefault(r["work_location"], []).append((_d(r["effective_from"]), _d(r["effective_to"]),
                                                         int(r["pay_zone"])))
    last = {}
    for r in _rows(target, "hris_job_history_2026-10-14.csv"):
        d = _d(r["effective_date"])
        if d <= BENCH:
            last[r["employee_id"]] = r          # rows run in date order within each worker
    grid = {(l, z): 0 for l in LV for z in (1, 2, 3)}
    for r in last.values():
        if r["worker_type"] != "employee" or r["status"] == "terminated" or not r["job_code"].startswith("DE"):
            continue
        z = [zz for a, b, zz in zones[r["work_location"]] if a <= BENCH and (b is None or BENCH <= b)][0]
        grid[(LV[int(r["job_code"][-1]) - 1], z)] += 1
    return grid


def open_reqs(target):
    seen, out = set(), {l: 0 for l in LV}
    for r in _rows(target, "ats_requisitions_2026-10-14.csv"):
        if r["requisition_id"] in seen or not r["job_code"].startswith("DE"):
            continue
        seen.add(r["requisition_id"])
        end = _d(r["status_date"]) if r["status"] in ("filled", "cancelled") else None
        if _d(r["opened_on"]) <= BENCH and (end is None or end > BENCH):
            out[LV[int(r["job_code"][-1]) - 1]] += 1
    return out


def main(target, golden):
    b, g, n = bands(target)
    hc, rq = headcount(target), open_reqs(target)
    wb = openpyxl.load_workbook(os.path.join(golden, "data_engineer_bands_2027.xlsx"), data_only=True)
    ws, hs, qs = wb["Bands 2027"], wb["Headcount 2026-10-01"], wb["Open requisitions"]
    bad = []
    for i, l in enumerate(LV):
        if ws.cell(row=5 + i, column=3).value != b[l][50]:
            bad.append("midpoint %s" % l)
        if ws.cell(row=5 + i, column=5).value != n[l]:
            bad.append("matched %s" % l)
        for j, z in enumerate((1, 2, 3)):
            if hs.cell(row=4 + i, column=2 + j).value != hc[(l, z)]:
                bad.append("headcount %s zone %d" % (l, z))
        if qs.cell(row=4 + i, column=2).value != rq[l]:
            bad.append("requisitions %s" % l)
    if (ws["B7"].value, ws["D7"].value) != (b["III"][25], b["III"][75]):
        bad.append("Data Engineer III quartiles")
    svg = open(os.path.join(golden, "data_engineer_bands_2027.svg"), encoding="utf-8").read()
    for l in LV:
        if "${:,}".format(b[l][50]) not in svg:
            bad.append("chart midpoint %s" % l)
    if "Data Engineer III midpoint ${:,}".format(b["III"][50]) not in svg:
        bad.append("chart title")
    if round(100 * g, 1) != 2.7:
        bad.append("movement %.3f" % g)
    if bad:
        raise SystemExit("verify: %d figures differ: %s" % (len(bad), ", ".join(bad)))
    print("verify: independent recompute agrees: %d band points, %d matched counts, %d headcount cells, %d "
          "requisition counts, the chart's labels, movement %.2f per cent a year" % (
              4 + 2, 4, 12, 4, 100 * g))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
