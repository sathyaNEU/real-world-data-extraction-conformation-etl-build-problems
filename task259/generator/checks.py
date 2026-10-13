"""The build's assertions: the input gates, the client's three rules executed (R1 to R3c), the semantic fork table
(SF1 to SF14), the ladder and every shortcut, the bins under every defensible estimator, the asks' truths and stops,
the separation line, the distractors, the realism sweep and the goldens against figures.json."""
from __future__ import annotations

import io
import json
import os
import re
import shutil
import tempfile

import numpy as np
import openpyxl
import pandas as pd

import analysis as A
import common as K
import load as L

FAILS, PASSES = [], []
GRADED = (("I", 50), ("II", 50), ("III", 25), ("III", 50), ("III", 75), ("IV", 50))


def ck(label, cond, detail=""):
    (PASSES if cond else FAILS).append(label if cond else "%s [%s]" % (label, detail))


def pdf_text(path):
    from pypdf import PdfReader
    return "\n".join(p.extract_text() for p in PdfReader(path).pages)


def xlsx_text(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    return "\n".join(str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row
                     if c is not None)


def doc_text(target, name):
    p = os.path.join(target, name)
    if name.endswith(".pdf"):
        return pdf_text(p)
    if name.endswith(".xlsx"):
        return xlsx_text(p)
    return open(p, encoding="utf-8").read()


def bin_margin(v):
    f = A.to_k(v)
    return min(v - (f - 500), (f + 500) - v)


def run(target, golden_dir, st, meta_path):
    w, fig = st["world"], st["fig"]
    f, G, R = fig["figures"], fig["G"], fig["R"]
    files = sorted(os.listdir(target))
    meta = json.load(open(meta_path))
    gold = f["call_DE_III_midpoint"]

    # input gates
    ck("gate: 10 or more files", len(files) >= 10, len(files))
    ck("gate: 3 or more formats", len({os.path.splitext(x)[1] for x in files}) >= 3)
    rows = {x: sum(1 for _ in open(os.path.join(target, x), encoding="utf-8")) - 1 for x in files if x.endswith(".csv")}
    ck("gate: one file of 25,000 or more rows (%s, %d)" % (K.EXTRACT, rows[K.EXTRACT]), rows[K.EXTRACT] >= 25_000)
    ck("gate: two or more distractors, shipped and named only in metadata.json",
       len(meta["distractor_files"]) >= 2 and all(x in files for x in meta["distractor_files"]) and
       not any("distract" in x for x in files))
    ck("gate: the goldens are the two deliverables the prompt names",
       sorted(os.listdir(golden_dir)) == sorted(K.DELIVERABLES))
    ck("gate: metadata lists every shipped file once as call, ask or distractor",
       sorted(meta["call_files"] + meta["ask_files"] + meta["distractor_files"]) == files)

    ex, pr, cw = L.extract(target), L.profiles(target), L.crosswalk(target)
    fam, focus = L.survey_match(target)
    df = G["df"]

    # the call
    ck("call: the Data Engineer III midpoint is filed at $170,000", gold == 170_000, gold)
    ck("call: the market's movement at a fixed level within employers is 2.7 per cent a year",
       f["movement_pct_a_year"] == 2.7)
    ck("call: across the pooled reports pay at a fixed level rises 1.1 per cent a year",
       f["pooled_movement_pct_a_year"] == 1.1)
    ck("call: the survey match is the data engineering family, primary focus data pipelines and platforms",
       (fam, focus) == ("Data Engineering", "data pipelines and platforms"))

    # R1: the field's raw material ships
    ck("R1: the extract is respondent microdata, one row per respondent, with filing dates and base pay",
       len(ex) == ex.respondent_id.nunique() and {"filed_on", "base_salary_usd", "employer_level"} <= set(ex.columns))
    cols = [c for x in files if x.endswith(".csv") for c in pd.read_csv(os.path.join(target, x), nrows=0).columns]
    ck("R1: no shipped column holds a movement, an aged pay, a percentile or a band point",
       not any(set(c.lower().split("_")) & {"aged", "ageing", "aging", "movement", "index", "band", "midpoint",
                                            "percentile", "trend", "growth"}
               for c in cols if c not in ("A_PCT10", "A_PCT25", "A_MEDIAN", "A_PCT75", "A_PCT90")))
    ck("R1: the golden reads %d matched reports filed over %d distinct days" % (len(df), df.filed_on.nunique()),
       len(df) > 4000 and df.filed_on.nunique() > 900)

    # R2: every Layer 1 item mishandled leaves the call where it was
    V = A.variants(ex, pr, cw, fam, focus)
    for tag in A.LAYER1:
        ck("R2 %s: the call unchanged (III midpoint %.0f)" % (tag, V[tag]["III"][50]),
           A.to_k(V[tag]["III"][50]) == gold)
    ck("R2: the call files only from the call files (distractors and ask files never read)",
       _call_from(target, meta["call_files"], fam, focus) == G["filed"])

    # R3a, R3b: the analyst test on the clean truth
    tb = w["tb"]
    m = tb[tb.matched_truth].copy()
    m["lv"] = [K.LEVELS[i] for i in m.level_idx]
    p = m[m.updated & ~m.promoted]
    g_pair = float(np.mean(np.log(p.base / p.prev_base) / p.gap))
    y, t = np.log(m.base.to_numpy()), -m.age.to_numpy()
    X = np.c_[pd.get_dummies(m.lv).to_numpy(float), t]
    g_pool = float(np.linalg.lstsq(X, y, rcond=None)[0][-1])
    emp_, lv_ = m.employer_canonical.to_numpy(), m.lv.to_numpy()
    yy, tt = A._demean(y, [emp_, lv_]), A._demean(t, [emp_, lv_])
    g_clean = float(tt @ yy / (tt @ tt))
    iii = (m.lv == "III").to_numpy()
    clean = {"R1": A.to_k(np.median(m.base.to_numpy()[iii])),
             "R2": A.to_k(np.median((m.base * np.exp(g_pair * m.age)).to_numpy()[iii])),
             "R3": A.to_k(np.median((m.base * np.exp(g_pool * m.age)).to_numpy()[iii])),
             "R4": A.to_k(np.median((m.base * np.exp(g_clean * m.age)).to_numpy()[iii]))}
    ck("R3a: the easy route (each report carried at the re-reporters' own growth, %.1f per cent a year) on the clean "
       "truth files %d, the named wrong answer" % (100 * g_pair, clean["R2"]),
       clean["R2"] == f["DE_III_paired_growth"] and clean["R2"] - gold >= 4000, clean)
    ck("R3a: the careful route (the level's movement read across the pooled reports, %.2f per cent a year) on the "
       "clean truth files %d, the second named wrong answer" % (100 * g_pool, clean["R3"]),
       clean["R3"] == f["DE_III_pooled_movement"] and gold - clean["R3"] >= 3000, clean)
    ck("R3b: only the decisive step changed (the movement read within employers, %.2f per cent a year) files the "
       "golden on the clean truth" % (100 * g_clean), clean["R4"] == gold, clean)
    ck("R3b: the lower rung's single correction (no carry) files %d on the clean truth" % clean["R1"],
       clean["R1"] == f["DE_III_unaged"] and gold - clean["R1"] >= 5000, clean)
    ck("R3b: the re-reporters' growth is the market's movement plus merit (%.3f against %.3f)" % (g_pair, g_clean),
       g_pair - g_clean > 0.015)
    ck("R3b: the pooled reports' movement falls short of the movement within employers by more than a point a year "
       "(%.3f against %.3f): the employers filing later pay less" % (g_pool, g_clean), g_clean - g_pool > 0.010)
    arr = _arrivals(df)
    ck("R3b: the share of matched reports from employers first seen after the window's first six months climbs from "
       "%.0f to %.0f per cent across the half-years, and they pay %.1f per cent less at a fixed level" % (
           100 * arr["share"][0], 100 * arr["share"][-1], 100 * arr["gap"]),
       arr["share"][0] < 0.05 and arr["share"][-1] > 0.35 and arr["gap"] > 0.06, arr)

    # R3c: no plant, no filed method
    texts = {x: doc_text(target, x) for x in (K.POLICY, K.METHODOLOGY, K.ARCHITECTURE, K.PEOPLE_NOTES)}
    every = "\n".join(texts.values())
    ck("R3c: no document states an ageing rule, a movement, a merit budget, an index or a prior band",
       not re.search(r"\bage[ds]?\b|\bageing\b|\baging\b|movement|merit|increase budget|prior band|\bindex|"
                     r"band of 20(2[0-5])|trend", every, re.I))
    ck("R3c: no document names the employers' arrival on the platform, its coverage, a mix or a fixed panel",
       not re.search(r"joined|arriv|coverage|employer mix|composition|same employers|fixed effect|panel|"
                     r"new employers|onboard", every, re.I))
    ck("R3c: no memo, draft, briefing or prior-year band file ships",
       not any(re.search(r"memo|draft|brief|bands?_20(2[0-6])", x) for x in files))
    ck("R3c: the policy files the estimand (who, when, which points, rounding) and nothing on carrying",
       all(s in texts[K.POLICY] for s in ("counts once", "1 October 2026", "nearest $1,000", "25th percentile")))

    # the semantic fork table
    ck("SF1 base pay: the extract's pay columns are annual base salary only",
       [c for c in ex.columns if "salary" in c or "pay" in c] == ["base_salary_usd", "previous_base_salary_usd"])
    ck("SF2 the market: policy 4.2 counts every matched respondent once (no weighting)",
       "Each matched respondent counts once" in texts[K.POLICY].replace("\n", " "))
    sh = _shortcuts(ex, pr, cw, fam, focus, G["g"], R["g_pair"], R["g_pooled"])
    ck("SF3 Tallyband's level for Tarnwick's: another figure (%s)" % sh["focus_L4_aged"],
       sh["focus_L4_aged"] != gold)
    ck("SF4 the data engineering family for the survey match: another figure (%s)" % sh["family_xwalk_aged"],
       sh["family_xwalk_aged"] != gold)
    last12 = df[[(K.BENCHMARK - d).days <= 365 for d in df.filed_on]]
    l12 = A.to_k(np.median(last12[last12.level == "III"].base_salary_usd))
    ck("SF5 only the last twelve months' reports, as filed: another figure (%d)" % l12, l12 != gold)
    ck("SF6 the crosswalk row in force today for every report: same figure",
       all(A.to_k(V["current_crosswalk"][lv][p_]) == G["filed"][lv][p_] for lv, p_ in GRADED))
    for sf, keys in (("SF7 percentile definitions", [k for k in A.DETERMINISM if k.startswith("pct_")]),
                     ("SF8 rounding each carried pay first", ["round_first"]),
                     ("SF9 the movement's estimator within employers", ["fe_employer_level", "fe_per_level",
                                                                         "panel_first_year", "panel_every_year"]),
                     ("SF10 the carrying form", ["linear_form", "compound_form"]),
                     ("SF11 whole months for days", ["months"]),
                     ("SF12 US-located only against US employers abroad", ["keep_nonus_only"]),
                     ("SF13 the employer as typed or tidied", ["fe_typed"])):
        ok = all(A.to_k(V[k][lv][p_]) == G["filed"][lv][p_] for k in keys for lv, p_ in GRADED)
        ck("%s: same figures for every graded point" % sf, ok,
           [(k, lv, p_, round(V[k][lv][p_])) for k in keys for lv, p_ in GRADED
            if A.to_k(V[k][lv][p_]) != G["filed"][lv][p_]])

    wide = {"every family's reports": _movement_wide(ex, "extract"), "the offers extract": _movement_wide(
        pd.read_csv(os.path.join(target, K.OFFERS)), "offers")}
    for k, gw in wide.items():
        mid = A.to_k(np.median(A.aged(df[df.level == "III"], gw)))
        ck("SF14 the movement within employers read from %s (%.2f per cent a year): the same call" % (k, 100 * gw),
           mid == gold, mid)

    # bins: every graded point mid-bin under every defensible estimator
    for lv, p_ in GRADED:
        vals = [V[k][lv][p_] for k in A.DETERMINISM] + [G["points"][lv][p_]]
        fk = G["filed"][lv][p_]
        mg = min(min(vals) - (fk - 500), (fk + 500) - max(vals))
        ck("bins: %s P%d (%d) holds a $%d margin under every estimator" % (lv, p_, fk, mg), mg >= 100, mg)
    ck("bins: the call sits %.0f from its nearest edge" % bin_margin(G["points"]["III"][50]),
       bin_margin(G["points"]["III"][50]) >= 250)

    # the ladder and every shortcut
    r = {"R0": A.to_k(R["R0"]["III"]), "R1": A.to_k(R["R1"]["III"][50]), "R3": A.to_k(R["R3"]["III"][50]),
         "R3_no_level": A.to_k(R["R3_no_level"]["III"][50]), "R4": A.to_k(R["R4"]["III"][50]),
         "R2": A.to_k(R["R2"]["III"][50]), "R2_all": A.to_k(R["R2_all"]["III"][50])}
    ck("ladder: R0 %d < R1 %d < R3 %d < R3 without level intercepts %d < golden %d < R2 %d < R2 on every pair %d" % (
        r["R0"], r["R1"], r["R3"], r["R3_no_level"], r["R4"], r["R2"], r["R2_all"]),
       r["R0"] < r["R1"] < r["R3"] < r["R3_no_level"] < r["R4"] < r["R2"] < r["R2_all"] and r["R4"] == gold)
    ck("ladder: every rung files its own figure %s" % r, len(set(r.values())) == 7)
    ck("ladder: the golden two rounding steps or more from every rung",
       all(abs(v - gold) >= 2000 for k, v in r.items() if k != "R4"), r)
    ck("ladder: the default path's paired growth is %.1f per cent a year over %d pairs" % (100 * R["g_pair"],
                                                                                           R["n_pair"]),
       f["paired_growth_pct_a_year"] == round(100 * R["g_pair"], 1) and R["n_pair"] > 800)
    ck("ladder: no shortcut lands on the golden %s" % sh, all(v != gold for k, v in sh.items()), sh)
    xs = {k[3:]: A.to_k(V[k]["III"][50]) for k in V if k.startswith("xs_")}
    ck("ladder: every cross-section reading of the movement files below the golden %s" % xs,
       all(v <= gold - 2000 for v in xs.values()), xs)
    fe_nl = A.to_k(np.median(A.aged(df[df.level == "III"], A.movement(df, "fe_no_level"))))
    ck("ladder: within employers but without an intercept for each level files %d" % fe_nl, fe_nl != gold, fe_nl)
    ck("ladder: every level moves the same way on each rung",
       all(R["R1"][lv][50] < R["R3"][lv][50] < G["points"][lv][50] < R["R2"][lv][50] for lv in A.LEVELS))
    for rung in ("R1", "R3", "R3_no_level", "R2"):
        offs = [(lv, q_) for lv, q_ in GRADED if (lv, q_) != ("III", 50) and
                A.to_k(R[rung][lv][q_]) == G["filed"][lv][q_]]
        ck("ladder: %s misses every Ask C figure" % rung, not offs, offs)

    # Ask A: the headcount grid
    hr, zn, req = L.hris(target), L.zones(target), L.ats(target)
    hc = fig["headcount"]
    tr = w["hris_truth"]
    tr = tr[(tr.family == "DE") & (tr.worker_type == "employee")]
    tgrid = tr.groupby(["level", "zone"]).size().unstack(fill_value=0).reindex(index=A.LEVELS, columns=[1, 2, 3],
                                                                                 fill_value=0)
    ck("Ask A: the golden grid equals the world's own count in every cell", (tgrid.values == hc.values).all())
    crack = A.headcount(hr, zn, zone_from="record")
    ck("Ask A cracker (record in force, the zone on the record): every cell off",
       (crack.values != hc.values).all(), crack.values.tolist())
    stops = {"active only": A.headcount(hr, zn, statuses=("active",)),
             "contingent kept": A.headcount(hr, zn, employees_only=False),
             "latest record (entered ahead read as current)": A.headcount(hr, zn, as_of=False)}
    for k, v in stops.items():
        ck("Ask A stop, %s: %d of 12 cells off" % (k, int((v.values != hc.values).sum())),
           (v.values != hc.values).sum() >= 3)
    rows_all = hr[hr.job_code.str.startswith("DE")]
    ck("Ask A stop, every row counted: %d rows for %d employees" % (len(rows_all), int(hc.values.sum())),
       len(rows_all) > 2 * hc.values.sum())
    ck("Ask A: every level and zone has employees, no level flat across its zones %s" % hc.values.tolist(),
       (hc.values > 0).all() and all(len(set(row)) > 1 for row in hc.values.tolist()))

    # Ask B: open requisitions
    tb_ = w["ats_truth"]
    tb_ = tb_[tb_.family == "DE"]
    tr_open = {lv: int(((tb_.level == lv) & (tb_.opened_on <= K.BENCHMARK) &
                        (tb_.ended_on.isna() | (tb_.ended_on > K.BENCHMARK))).sum()) for lv in A.LEVELS}
    ck("Ask B: the golden count equals the world's own at every level", tr_open == fig["open_reqs"], tr_open)
    crb = A.open_requisitions(req, end_from="closed_on")
    ck("Ask B cracker (distinct requisitions, closed_on as the end): every level off %s" % crb,
       all(crb[lv] != fig["open_reqs"][lv] for lv in A.LEVELS))
    st_open = req.drop_duplicates("requisition_id")
    st_open = st_open[st_open.job_code.str.startswith("DE") & (st_open.status == "open")]
    so = {lv: int((st_open.job_code.str[-1].map(A.ROMAN_OF) == lv).sum()) for lv in A.LEVELS}
    ck("Ask B stop, status open as exported: %d of 4 levels off %s" % (
        sum(so[lv] != fig["open_reqs"][lv] for lv in A.LEVELS), so),
       sum(so[lv] != fig["open_reqs"][lv] for lv in A.LEVELS) >= 3)
    pr_rows = req[req.job_code.str.startswith("DE")]
    pr_rows = pr_rows[[o <= K.BENCHMARK and (c is None or c > K.BENCHMARK) for o, c in
                       zip(pr_rows.opened_on, pr_rows.closed_on)]]
    pc = {lv: int((pr_rows.job_code.str[-1].map(A.ROMAN_OF) == lv).sum()) for lv in A.LEVELS}
    ck("Ask B stop, posting rows counted: every level off %s" % pc,
       all(pc[lv] != fig["open_reqs"][lv] for lv in A.LEVELS))

    # separation
    main_docs = "\n".join(texts[x] for x in (K.POLICY, K.METHODOLOGY, K.ARCHITECTURE))
    ck("separation: no main-call document names an ask file, a requisition, a leave, a contingent worker or the "
       "zone change", not re.search(r"requisition|hris|\bats\b|leave|contingent|1 July 2026|zone map", main_docs,
                                     re.I))
    ck("separation: the asks never read a call file or a distractor",
       _asks_from(target, meta["ask_files"]) == (hc.values.tolist(), fig["open_reqs"]))

    # realism sweep
    for x in files:
        ck("realism: no em dash in %s" % x, "\u2014" not in doc_text(target, x))
    for x in K.DELIVERABLES:
        txt = xlsx_text(os.path.join(golden_dir, x)) if x.endswith(".xlsx") else open(
            os.path.join(golden_dir, x), encoding="utf-8").read()
        ck("realism: no em dash in the golden %s" % x, "\u2014" not in txt)
    ck("realism: filing dates inside the 36 months to the extract, previous reports earlier",
       ex.filed_on.min() >= K.EXTRACT_FROM and ex.filed_on.max() <= K.BENCHMARK and
       all(pf < fd for pf, fd in zip(ex.previous_filed_on, ex.filed_on) if pf is not None))
    ck("realism: HRIS rows past the export are changes entered ahead, none past 16 November 2026",
       max(hr.effective_date) <= dt_date(2026, 11, 16) and (hr.effective_date > K.HR_EXPORT).sum() > 0)
    ck("realism: ATS dates on or before the export",
       all(d <= K.HR_EXPORT for c in ("opened_on", "closed_on", "status_date") for d in req[c] if d is not None))
    ck("realism: no record dated on the benchmark day in the HRIS or ATS (no boundary reading)",
       K.BENCHMARK not in set(hr.effective_date) and K.BENCHMARK not in set(req.opened_on) and
       K.BENCHMARK not in set(d for d in req.status_date if d is not None))
    ck("realism: matched counts differ by level %s" % G["n"], len(set(G["n"].values())) == 4)
    ck("realism: the policy PDF and the workbook carry no library banner",
       b"ReportLab" not in open(os.path.join(target, K.POLICY), "rb").read() and
       b"openpyxl" not in open(os.path.join(target, K.ARCHITECTURE), "rb").read())
    ck("realism: personas from guard.py names only", all(n in every + _golden_text(golden_dir)
                                                         for n in K.PEOPLE.values()))

    # goldens against figures.json
    wb = openpyxl.load_workbook(os.path.join(golden_dir, K.DELIVERABLES[0]), data_only=True)
    bands = wb["Bands 2027"]
    ck("golden workbook: the four midpoints and Data Engineer III's quartiles as figures.json",
       [bands.cell(row=5 + i, column=3).value for i in range(4)] == [f["midpoint_%s" % lv] for lv in A.LEVELS] and
       (bands["B7"].value, bands["D7"].value) == (f["DE_III_p25"], f["DE_III_p75"]))
    hcs = wb["Headcount 2026-10-01"]
    ck("golden workbook: the headcount grid as figures.json",
       [[hcs.cell(row=4 + i, column=2 + j).value for j in range(3)] for i in range(4)] == hc.values.tolist())
    rqs = wb["Open requisitions"]
    ck("golden workbook: the open requisitions as figures.json",
       [rqs.cell(row=4 + i, column=2).value for i in range(4)] == [f["open_requisitions_%s" % lv] for lv in A.LEVELS])
    svg = open(os.path.join(golden_dir, K.DELIVERABLES[1]), encoding="utf-8").read()
    ck("golden chart: every midpoint and Data Engineer III's quartiles labelled as figures.json",
       all("${:,}".format(f["midpoint_%s" % lv]) in svg for lv in A.LEVELS) and
       "${:,}".format(f["DE_III_p25"]) in svg and "${:,}".format(f["DE_III_p75"]) in svg)

    print("checks: %d passed, %d failed" % (len(PASSES), len(FAILS)))
    if FAILS:
        for x in FAILS:
            print("FAIL", x)
        raise SystemExit(1)


def dt_date(y, m, d):
    import datetime
    return datetime.date(y, m, d)


def _golden_text(golden_dir):
    return xlsx_text(os.path.join(golden_dir, K.DELIVERABLES[0]))


def _copy(target, names):
    tmp = tempfile.mkdtemp(dir=os.path.dirname(os.path.abspath(target)), suffix=".chk")
    for x in names:
        shutil.copy(os.path.join(target, x), tmp)
    return tmp


def _call_from(target, names, fam, focus):
    tmp = _copy(target, names)
    try:
        return A.golden(L.extract(tmp), L.profiles(tmp), L.crosswalk(tmp), fam, focus)["filed"]
    finally:
        shutil.rmtree(tmp)


def _asks_from(target, names):
    tmp = _copy(target, names)
    try:
        return (A.headcount(L.hris(tmp), L.zones(tmp)).values.tolist(), A.open_requisitions(L.ats(tmp)))
    finally:
        shutil.rmtree(tmp)


def _arrivals(df):
    """Each employer's first report (current or previous) in the matched reports; the share of reports by half-year
    from employers first seen inside the window, and how much less they pay within level, carried at the golden."""
    first = {}
    for e, d, pd_ in zip(df.employer.map(A.norm_employer), df.filed_on, df.previous_filed_on):
        for x in (d, pd_):
            if x is not None and (e not in first or x < first[e]):
                first[e] = x
    cutoff = K.EXTRACT_FROM + (K.BENCHMARK - K.EXTRACT_FROM) / 6
    late = np.array([first[e] > cutoff for e in df.employer.map(A.norm_employer)])
    half = (A.age_years(df.filed_on) // 0.5).astype(int)
    share = [float(late[half == h].mean()) for h in sorted(set(half), reverse=True)]
    lp = np.log(A.aged(df, A.movement(df, "fe")))
    lv = df.level.to_numpy()
    resid = lp - pd.Series(lp).groupby(lv).transform("mean").to_numpy()
    return dict(share=share, gap=float(resid[~late].mean() - resid[late].mean()))


def _movement_wide(frame, kind):
    """The movement within employers read from a wider body of reports: every family's US full-time reports in the
    extract (family and Tallyband level intercepts), or the offers extract."""
    if kind == "extract":
        d = frame[(frame.country == "US") & (frame.employment_type == "full_time")]
        y, dates = np.log(d.base_salary_usd.to_numpy(float)), d.filed_on
        strata = [d.tallyband_level.to_numpy(), d.job_family.to_numpy()]
    else:
        d = frame
        y = np.log(d.base_offered_usd.to_numpy(float))
        dates = [L.as_date(x) for x in d.offered_on]
        strata = [d.tallyband_level.to_numpy(), d.family.to_numpy()]
    t = -A.age_years(list(dates))
    emp = d.employer.map(A.norm_employer).to_numpy()
    yy, tt = A._demean(y, strata + [emp]), A._demean(t, strata + [emp])
    return float(tt @ yy / (tt @ tt))


def _shortcuts(ex, pr, cw, fam, focus, g, g_pair, g_pool):
    d = ex.merge(pr[["respondent_id", "primary_focus"]], on="respondent_id")
    d["level"] = A.map_levels(d, cw)
    usft = (d.country == "US") & (d.employment_type == "full_time")
    sel = {"family_L4": (d.job_family == fam) & (d.tallyband_level == "L4") & usft,
           "focus_L4": (d.job_family == fam) & (d.primary_focus == focus) & (d.tallyband_level == "L4") & usft,
           "family_xwalk": (d.job_family == fam) & (d.level == "III") & usft}
    out = {}
    for k, msk in sel.items():
        x = d[msk]
        out[k + "_unaged"] = A.to_k(np.median(x.base_salary_usd))
        out[k + "_aged"] = A.to_k(np.median(A.aged(x, g)))
        out[k + "_paired"] = A.to_k(np.median(A.aged(x, g_pair)))
        out[k + "_pooled"] = A.to_k(np.median(A.aged(x, g_pool)))
    return out
