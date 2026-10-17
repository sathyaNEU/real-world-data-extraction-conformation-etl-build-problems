"""The pack's documents: the historical society's transcription of the water-works reports, the State Water
Resources Board's design-drought guidance, the town's drought management plan, its withdrawal permit, the gauge's
station description, Castlewick's offer, the Board's regional regression report (a distractor), the meter register
workbook, the data dictionary and the folder's contents list."""
import datetime as dt

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import letter                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import inch                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import numpy as np    # noqa: E402

import pack as P      # noqa: E402
import sim as S       # noqa: E402
import world as W     # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("LiberationSerif", "Regular"), ("LiberationSerif-Bold", "Bold"), ("LiberationSerif-Italic", "Italic"),
              ("LiberationSans", "Regular"), ("LiberationSans-Bold", "Bold")):
    try:
        pdfmetrics.getFont(nm)
    except KeyError:
        pdfmetrics.registerFont(TTFont(nm, FONT_DIR + nm.split("-")[0] + "-" + f + ".ttf"))
INK = colors.HexColor("#1b1d1f")
GREY = colors.HexColor("#5b6067")
NAVY = colors.HexColor("#1f3a5f")
TEAL = colors.HexColor("#1d5b63")
BROWN = colors.HexColor("#5a3e2b")
SLATE = colors.HexColor("#3d4f5c")
BODY = ParagraphStyle("body", fontName="LiberationSerif", fontSize=10.5, leading=14.2, textColor=INK, spaceAfter=6)
SANS = ParagraphStyle("sans", fontName="LiberationSans", fontSize=9.6, leading=13.2, textColor=INK, spaceAfter=6)
QUOTE = ParagraphStyle("quote", fontName="LiberationSerif-Italic", fontSize=10, leading=13.6, textColor=INK,
                       leftIndent=22, rightIndent=22, spaceAfter=8)
CELL = ParagraphStyle("cell", fontName="LiberationSans", fontSize=8.4, leading=10.4, textColor=INK)
CELLB = ParagraphStyle("cellb", parent=CELL, fontName="LiberationSans-Bold")
SCELL = ParagraphStyle("scell", fontName="LiberationSerif", fontSize=8.8, leading=10.8, textColor=INK)
SCELLB = ParagraphStyle("scellb", parent=SCELL, fontName="LiberationSerif-Bold")
SMALL = ParagraphStyle("small", fontName="LiberationSans", fontSize=8.3, leading=10.8, textColor=GREY)


def styles(accent, serif=False):
    face = "LiberationSerif" if serif else "LiberationSans"
    h1 = ParagraphStyle("h1", fontName=face + "-Bold", fontSize=14, leading=18, textColor=accent, spaceAfter=2)
    h1s = ParagraphStyle("h1s", fontName=face, fontSize=9.3, leading=12.5, textColor=GREY, spaceAfter=10)
    h2 = ParagraphStyle("h2", fontName=face + "-Bold", fontSize=10.5, leading=14, textColor=INK, spaceBefore=7,
                        spaceAfter=3)
    return h1, h1s, h2


def _doc(path, title, author, footer, layout="left"):
    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(GREY)
        if layout == "left":
            canvas.setFont("LiberationSans", 7.5)
            canvas.drawString(0.9 * inch, 0.6 * inch, footer)
            canvas.drawRightString(letter[0] - 0.9 * inch, 0.6 * inch, "Page %d" % doc.page)
        elif layout == "centre":
            canvas.setFont("LiberationSerif", 8.5)
            canvas.drawCentredString(letter[0] / 2, 0.6 * inch, "%s  -  %d" % (footer, doc.page))
        elif layout == "header":
            canvas.setFont("LiberationSans", 8)
            canvas.drawString(0.9 * inch, letter[1] - 0.55 * inch, footer)
            canvas.setStrokeColor(GREY)
            canvas.setLineWidth(0.5)
            canvas.line(0.9 * inch, letter[1] - 0.6 * inch, letter[0] - 0.9 * inch, letter[1] - 0.6 * inch)
            canvas.drawCentredString(letter[0] / 2, 0.55 * inch, str(doc.page))
        else:
            canvas.setFont("LiberationSans", 7)
            canvas.drawRightString(letter[0] - 0.9 * inch, 0.6 * inch, "%s  |  page %d" % (footer, doc.page))
        canvas.restoreState()
    d = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.9 * inch, rightMargin=0.9 * inch, topMargin=0.8 * inch,
                          bottomMargin=0.9 * inch, title=title, author=author, subject="", creator="", producer="")
    return d, frame


def _table(rows, widths, head="#e6eef0", serif=False):
    c, cb = (SCELL, SCELLB) if serif else (CELL, CELLB)
    data = [[Paragraph(str(v), cb if i == 0 else c) for v in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa0a6")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(head)),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 4),
                           ("RIGHTPADDING", (0, 0), (-1, -1), 4)]))
    return t


def long_date(d):
    if isinstance(d, str):
        d = dt.date.fromisoformat(d)
    return d.strftime("%B ") + str(d.day) + d.strftime(", %Y")


MONTHS = ["Jan.", "Feb.", "March", "April", "May", "June", "July", "Aug.", "Sept.", "Oct.", "Nov.", "Dec."]


def old_date(d):
    return "%s %d" % (MONTHS[d.month - 1], d.day)


# ---------------------------------------------------------------- the water-works reports, transcribed
AGREEMENT = ("... and the said Commissioners shall cause the pumps of the Water Works to be stopped whenever the flow "
             "of the river over the dam, taken as the average of any seven consecutive days, shall be less than two "
             "millions of gallons in twenty-four hours, and shall not start them again until the said flow shall have "
             "risen above that quantity; and the Company shall have the use of all water passing the dam.")
REMARKS_DRY = ["Supply drawn from the Brook Street reservoir while the pumps were stopped.",
               "Mill on short time through September.", "Fire department notified of reduced pressure.",
               "Residents asked to stop lawn sprinkling.", "Reservoir at its lowest since the works opened.",
               "Water carted to the hill district for nine days.", "Brook Street reservoir drawn to within four feet.",
               "Street sprinkling suspended.", "Watering troughs on Main Street shut off.",
               "Leaks on the Elm Street main repaired to save water.", "Weir read morning and evening while the "
               "pumps were stopped.", "Hydrants on the Mill Road closed until the pumps started.",
               "Notice printed asking householders to save water.", "Brook Street reservoir down eleven feet.",
               "Mill pond drawn below the flashboards.", "Laundries asked to work alternate days."]
TWICE = ["Pumping resumed after rain, stopped again.", "Pumps started after the rain and stopped a second time.",
         "A few days' pumping after rain.", "Pumps run again after a heavy rain, then stopped.",
         "Started after rain; flow fell off and pumps stopped again.", "Rain gave a short run of pumping.",
         "Pumps restarted when the river rose, stopped again within the week."]
REMARKS_WET = ["New filter bed completed.", "Service pipes renewed on Main Street.", "Pump No. 2 rebuilt.",
               "Freshet in April damaged the intake screen.", "Extension of mains to the Mill Road district.",
               "Meters placed on all business services.", "High water in March; no damage.",
               "Second boiler installed at the pumping station.", "Brook Street reservoir cleaned.",
               "Intake screen replaced.", "Hydrants added on Pleasant Street.", "Pump No. 1 rebuilt.",
               "New coal shed built at the pumping station.", "Main laid to the cemetery.", "Gate house roof repaired.",
               "Flood in November; dam inspected.", "Ice cleared from the intake in February.",
               "Valves renewed on the Brook Street main."]


def report_table(x):
    rows = [["Year", "Pumps stopped under the mill agreement", "Lowest seven-day flow over the dam, as reported",
             "Remarks"]]
    order = np.random.default_rng([W.SEED, 81])
    dry = [REMARKS_DRY[i] for i in order.permutation(len(REMARKS_DRY))]
    twice = [TWICE[i] for i in order.permutation(len(TWICE))]
    wet = [REMARKS_WET[i] for i in order.permutation(len(REMARKS_WET))]
    for y in range(W.FIRST_REPORT, W.LAST_REPORT + 1):
        if y in W.MISSING_REPORTS:
            rows.append([str(y), "Report not found", "", ""])
            continue
        s = x["seasons"][y]
        rng = np.random.default_rng([W.SEED, 80, y])
        if not s["suspended"]:
            rem = wet.pop() if rng.random() < 0.35 and wet else ""
            rows.append([str(y), "No stoppage", "", rem])
            continue
        spans = "; ".join("%s to %s" % (old_date(a), old_date(b)) for a, b in s["spans"])
        rep = S.reported_flow(y, s["flow"])
        if y < W.UNITS_CHANGE:
            shown = "{:,} gal.".format(rep)
        else:
            shown = "%.2f sec.-ft." % rep
        rem = dry.pop() if rng.random() < 0.6 and dry else ""
        if len(s["spans"]) > 1:
            rem = (twice.pop() + " " + rem).strip()
        rows.append([str(y), spans, shown, rem])
    return rows


def write_reports(path, x):
    H1, H1S, H2 = styles(BROWN, serif=True)
    doc, frame = _doc(path, "Quillfield Water Works: Pumping Stopped under the Mill Agreement, 1894 to 1950",
                      W.PEOPLE["archivist"], "%s  |  Water Works reports, 1894 to 1950" % W.SOCIETY, layout="centre")
    s = [Paragraph("Quillfield Water Works: Pumping Stopped under the Mill Agreement, 1894 to 1950", H1),
         Paragraph("From the Annual Reports of the Board of Water Commissioners of the Town of Quillfield. Transcribed "
                   "by %s, %s, %s. Collection QHS-114." % (W.PEOPLE["archivist"], W.TITLES["archivist"].split(",")[0],
                                                            W.SOCIETY), H1S),
         Paragraph("The Water Works", H2),
         Paragraph("The town built its water works in 1893 at the dam of the %s on the %s, pumping from the mill pond "
                   "to the Brook Street reservoir. The company held the water rights at the dam, and the town's right to "
                   "pump was granted by an agreement of June 12, 1893, which bound the Water Commissioners as "
                   "follows:" % (W.MILL, W.RIVER), BODY),
         Paragraph(AGREEMENT, QUOTE),
         Paragraph("In each season in which the pumps were stopped under the agreement, the Superintendent gauged the "
                   "flow over the spillway every day with the weir the company kept at the dam, and his report to the "
                   "Commissioners gave the dates the pumps were stopped and the lowest average flow of seven consecutive "
                   "days. Until 1920 the reports give that flow in gallons a day; from 1921, when the State Engineer's "
                   "office began to collect the gaugings, in second-feet (cubic feet per second). Flows are transcribed "
                   "as printed.", BODY),
         Paragraph("The reports for 1909, 1918 and 1944 have not been found in the town's records or the Society's. "
                   "The mill closed in 1950, ending the agreement, and the State Hydrologic Survey began recording at "
                   "the dam in April 1951.", BODY),
         Paragraph("The Reports", H2),
         _table(report_table(x), [0.5 * inch, 2.2 * inch, 1.55 * inch, 2.45 * inch], head="#efe6dc", serif=True),
         Spacer(1, 6),
         Paragraph("Transcribed from the bound reports in the Quillfield Public Library and the Society's collection. "
                   "Spelling of place names modernised; figures as printed.", SMALL)]
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------- the Board's guidance
def write_guidance(path):
    H1, H1S, H2 = styles(NAVY)
    doc, frame = _doc(path, "Guidance %s: Design Drought Flows for Small Public Water Systems" % W.GUIDANCE_NO, W.BOARD,
                      "%s  |  Guidance %s (revised March 2023)" % (W.BOARD, W.GUIDANCE_NO), layout="header")
    s = [Paragraph("Guidance %s: Design Drought Flows for Small Public Water Systems" % W.GUIDANCE_NO, H1),
         Paragraph("%s. Revised March 2023. Applies to public water systems serving fewer than 25,000 people that "
                   "withdraw from a river or stream." % W.BOARD, H1S),
         Paragraph("1. Purpose", H2),
         Paragraph("A system that draws from a stream plans its drought supply against a design drought flow at its "
                   "intake. This guidance sets how that flow is defined and estimated for drought management plans, "
                   "firm-capacity statements and emergency supply arrangements filed with the Board.", SANS),
         Paragraph("2. Definitions", H2),
         Paragraph("<b>Climatic year.</b> The twelve months from April 1 to the following March 31, named by the "
                   "calendar year in which it begins.", SANS),
         Paragraph("<b>7-day low flow.</b> For a climatic year, the lowest mean of seven consecutive daily mean flows "
                   "whose last day falls in that year.", SANS),
         Paragraph("<b>Design drought flow.</b> The 7-day low flow with a 50-year recurrence interval: the flow below "
                   "which a climatic year's 7-day low flow falls with a probability of 0.02.", SANS),
         Paragraph("3. Estimation", H2),
         Paragraph("The design drought flow is estimated from a log-Pearson Type III distribution of the climatic-year "
                   "7-day low flows. The skew coefficient of their logarithms is taken as the Board's regional low-flow "
                   "skew, -0.30; the mean and standard deviation of the logarithms are estimated by maximum likelihood "
                   "from the river's record at the intake: the complete climatic years of the approved record of a "
                   "streamgage at or near the intake and, for earlier years, the records of the river kept at the "
                   "site. Flows at a gage away from the intake are transferred by drainage area.", SANS),
         Paragraph("Flows at other recurrence intervals are read from the same fitted distribution; the 10-year 7-day "
                   "low flow is the flow with a probability of 0.10.", SANS),
         Paragraph("4. Filing", H2),
         Paragraph("The plan states the design drought flow in cubic feet per second, the record used and the "
                   "withdrawal the system's permit allows at that flow. The Board reviews filed estimates when a plan "
                   "is revised or an emergency supply arrangement is made.", SANS)]
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------- the town's drought management plan
def write_plan(path):
    H1, H1S, H2 = styles(TEAL)
    doc, frame = _doc(path, "Town of Quillfield Drought Management Plan", "Quillfield Water Department",
                      "Town of Quillfield Water Department  |  Drought Management Plan, 2024 revision")
    dem = [["Year", "Stage 2 demand, MGD", "Average day, MGD", "Maximum day, MGD"]]
    for y, avg, peak in ((2025, "0.86", "1.41"), (2026, "0.87", "1.43"), (2027, "0.88", "1.45"), (2028, "0.89", "1.47")):
        dem.append([str(y), "%.2f" % W.STAGE2_DEMAND[y], avg, peak])
    s = [Paragraph("Drought Management Plan", H1),
         Paragraph("Town of Quillfield Water Department. 2024 revision, adopted by the Select Board on May 14, 2024.", H1S),
         Paragraph("1. Supply", H2),
         Paragraph("The town's water comes from the %s at the old mill dam, treated at the Mill Street plant, and from "
                   "the %s, which pumps directly into the distribution system. The plant also delivers wholesale "
                   "water to the %s under its 2011 agreement. Withdrawals from the river are governed by permit %s."
                   % (W.RIVER, W.WELLS, W.DISTRICT, W.PERMIT_NO), SANS),
         Paragraph("2. Drought stages", H2),
         Paragraph("The Select Board declares and rescinds stages by public notice on the advice of the Water "
                   "Superintendent, using the State Hydrologic Survey's gauge at the dam. Stage 1 (voluntary "
                   "conservation) applies when the 7-day mean flow falls below 6.0 cubic feet per second between June 1 "
                   "and November 30; Stage 2 (mandatory restrictions on outdoor use, as the permit requires) when it "
                   "falls below 4.0. A stage is rescinded once the 7-day mean has stood at or above its level for 14 "
                   "days. A notice takes effect the day after it is posted; a rescission keeps the stage in effect "
                   "through the day it is posted.", SANS),
         Paragraph("3. Stage 2 demand", H2),
         Paragraph("The demand the system must meet under Stage 2 restrictions, from the department's 2024 demand "
                   "projection:", SANS),
         _table(dem, [0.8 * inch, 1.5 * inch, 1.4 * inch, 1.4 * inch], head="#e2efef"),
         Spacer(1, 6),
         Paragraph("4. Emergency interconnect", H2),
         Paragraph("The town will hold firm emergency capacity on an interconnect with a neighbouring system, reserved "
                   "each contract year as the smallest block offered that covers the Stage 2 demand less the "
                   "withdrawal the permit allows at the design drought flow of the State Water Resources Board's "
                   "Guidance %s." % W.GUIDANCE_NO, SANS),
         Paragraph("5. Review", H2),
         Paragraph("The Water Superintendent reviews this plan after any Stage 2 period and before each interconnect "
                   "contract year.", SANS)]
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------- the withdrawal permit
def write_permit(path):
    H1, H1S, H2 = styles(SLATE)
    doc, frame = _doc(path, "Water Withdrawal Permit %s" % W.PERMIT_NO, W.BOARD,
                      "%s  |  Permit %s" % (W.BOARD, W.PERMIT_NO), layout="vendor")
    s = [Paragraph("Water Withdrawal Permit %s" % W.PERMIT_NO, H1),
         Paragraph("Issued by the %s to the Town of Quillfield, March 3, 2021; expires March 2, 2031." % W.BOARD, H1S),
         Paragraph("Condition 1. Location", H2),
         Paragraph("Withdrawal from the %s at the Mill Street intake on the mill dam, drainage area %.1f square "
                   "miles." % (W.RIVER, W.DRAINAGE_SQMI), SANS),
         Paragraph("Condition 2. Quantity", H2),
         Paragraph("Withdrawals shall not exceed 1.40 million gallons a day. The permittee shall pass at least 0.50 "
                   "cubic feet per second over the dam at all times, and shall withdraw only the river's flow above that "
                   "passby.", SANS),
         Paragraph("Condition 3. Drought", H2),
         Paragraph("The permittee shall maintain a drought management plan approved by the Board and shall impose "
                   "mandatory restrictions on outdoor water use whenever the 7-day mean flow at the State Hydrologic "
                   "Survey gauge %s falls below 4.0 cubic feet per second between June 1 and November 30." % W.STATION_NO,
                   SANS),
         Paragraph("Condition 4. Reporting", H2),
         Paragraph("Daily withdrawals shall be metered and reported to the Board monthly.", SANS)]
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------- the station description
def write_station(path):
    H1, H1S, H2 = styles(NAVY)
    doc, frame = _doc(path, "Station Description %s %s" % (W.STATION_NO, W.STATION), W.SURVEY,
                      "%s  |  Station %s" % (W.SURVEY, W.STATION_NO), layout="header")
    rows = [["Item", "Description"],
            ["Location", "Left bank, 300 feet upstream from the Quillfield mill dam and the town's Mill Street intake."],
            ["Drainage area", "%.1f square miles." % W.DRAINAGE_SQMI],
            ["Period of record", "April 1951 to current year."],
            ["Gage", "Water-stage recorder with satellite telemetry; datum of gage 412.30 feet above sea level. "
                     "Fifteen-minute record from 1989; daily means computed on local standard time."],
            ["Control", "Low-water control is the bedrock riffle 60 feet below the gage."],
            ["Regulation and diversions", "No regulation upstream. The town's withdrawal is made at the dam, below the "
                                          "gage."],
            ["Record accuracy", "Records good, except estimated daily values during ice, which are fair."],
            ["Approval status", "Records approved through September 30, 2025. Later records are provisional."]]
    s = [Paragraph("Station Description: %s, %s" % (W.STATION_NO, W.STATION), H1),
         Paragraph("%s, Water Data Section. Revised November 2025." % W.SURVEY, H1S),
         _table(rows, [1.5 * inch, 5.1 * inch], head="#e6eaf0")]
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------- Castlewick's offer
def write_offer(path):
    H1, H1S, H2 = styles(SLATE, serif=True)
    doc, frame = _doc(path, "Firm Interconnect Capacity, 2027 to 2031", W.PEOPLE["castlewick"],
                      "City of Castlewick Department of Utilities", layout="centre")
    s = [Paragraph("City of Castlewick, Department of Utilities", H1),
         Paragraph("October 6, 2026", H1S),
         Paragraph("Mr. Jaime Galloway<br/>Water Superintendent<br/>Town of Quillfield", BODY),
         Paragraph("Re: Firm capacity on the Castlewick and Quillfield interconnect, 2027 to 2031", BODY),
         Paragraph("Dear Mr. Galloway,", BODY),
         Paragraph("Following our meeting in September, the City offers firm capacity on the Route 9 interconnect for "
                   "the five contract years 2027 to 2031 on these terms:", BODY),
         Paragraph("1. Firm capacity is reserved in blocks of 0.1 million gallons a day, up to 0.8 million gallons a day, "
                   "the interconnect's hydraulic limit. The Town names its blocks for each contract year by December 31 "
                   "of the year before.", BODY),
         Paragraph("2. The reservation charge is $28,400 a block a year. Water taken is billed at $4.12 per thousand "
                   "gallons.", BODY),
         Paragraph("3. Reserved capacity is available on demand, including during the City's own Stage 1 and Stage 2 "
                   "drought conditions.", BODY),
         Paragraph("4. Water taken above the reserved blocks is supplied only as the City's operations allow.", BODY),
         Paragraph("We would be glad to bring the agreement to our Council in January once the Select Board has "
                   "acted.", BODY),
         Paragraph("Sincerely,<br/><br/>%s<br/>%s" % (W.PEOPLE["castlewick"], W.TITLES["castlewick"]), BODY)]
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------- the Board's regional regression report (distractor)
REG_GAUGES = [("01-1405.10", "Tarrant Brook near Copperfield", 18.2, 12), ("01-1422.00", "Kettle River at Fenwick", 63.4, 31),
              ("01-1460.30", "Sorrel Brook at Danby Corner", 9.7, 8), ("01-1488.00", "Wending River near Harlow", 112.0, 24),
              ("01-1513.50", "Pike Brook at Upper Orrin", 27.5, 19), ("01-1540.00", "Sallow River at Cheswick", 88.1, 36),
              ("01-1602.20", "Coldstream Brook near Lisle", 14.9, 15), ("01-1655.00", "Blackmere River at Tolland Falls", 204.0, 29),
              ("01-1701.40", "Quarry Brook near Ellery", 6.3, 6), ("01-1730.00", "Hatch River at Moreland", 47.7, 22),
              ("01-1786.60", "Pickerel Brook at North Penry", 21.1, 41), ("01-1822.00", "Elverly River near Hartsbury", 156.0, 17),
              ("01-1904.00", "Brannagh Brook at Wells Corner", 11.8, 11), ("01-1977.50", "Lantern River at Coyle", 72.9, 28),
              ("01-2033.00", "Fellow Brook near Mount Ardell", 33.6, 14), ("01-2105.40", "Corran River at East Tilby", 128.0, 33),
              ("01-2168.00", "Minnow Brook at Pound Hill", 8.9, 9), ("01-2240.20", "Thistle River near Oakhurst", 95.2, 26)]


def write_regression(path):
    H1, H1S, H2 = styles(NAVY)
    doc, frame = _doc(path, "Regional Equations for Low-Flow Statistics at Ungaged Sites", W.BOARD,
                      "%s  |  Technical Report 2019-3" % W.BOARD, layout="header")
    rows = [["Station", "Name", "Drainage area, sq mi", "Stratified drift, %", "Years of record"]]
    for no, nm, da, sd in REG_GAUGES:
        rng = np.random.default_rng([W.SEED, 90, int(da * 10)])
        rows.append([no, nm, "%.1f" % da, "%d" % sd, "%d" % int(rng.integers(24, 71))])
    s = [Paragraph("Regional Equations for Low-Flow Statistics at Ungaged Sites", H1),
         Paragraph("%s, Technical Report 2019-3. Hydrologic Region 2 (central uplands)." % W.BOARD, H1S),
         Paragraph("Equations", H2),
         Paragraph("For ungaged sites in Region 2 the 7-day, 10-year and 7-day, 2-year low flows are estimated from the "
                   "drainage area (A, square miles) and the percentage of the basin in stratified drift (SD):", SANS),
         Paragraph("log10(7Q10) = -1.402 + 1.031 log10(A) + 0.0182 SD; standard error of prediction 0.27 log units.", SANS),
         Paragraph("log10(7Q2) = -0.951 + 1.008 log10(A) + 0.0149 SD; standard error of prediction 0.19 log units.", SANS),
         Paragraph("The equations apply to basins of 5 to 210 square miles with 3 to 45 percent stratified drift and no "
                   "significant regulation. They are not used at gaged sites, where the gage's record governs.", SANS),
         Paragraph("Gages used", H2),
         _table(rows, [0.95 * inch, 2.35 * inch, 1.1 * inch, 1.05 * inch, 0.9 * inch], head="#e6eaf0")]
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------- the meter register workbook
METERS = [
    ("WTP-RAW", "Raw water venturi", "Intake pump station, Mill Street", "River water drawn at the intake", "MG", "2009-05-18"),
    ("WTP-FIN", "Finished water magnetic meter", "Mill Street plant, clearwell effluent to the high-service header",
     "All finished water leaving the plant", "MG", "2009-05-18"),
    ("EH-01", "Elm Hollow master meter", "Mill Street plant, high-service header, wholesale connection",
     "Wholesale water to the Elm Hollow Water District", "kgal", "2011-07-01"),
    ("HR-WF", "Wellfield discharge totalizer", "Hollow Road wellfield, combined discharge to the distribution system",
     "Water pumped by Hollow Road wells 1 and 2", "gal", "2010-06-08"),
    ("BR-01", "Brook Street reservoir level", "Brook Street reservoir", "Storage level, feet above floor", "ft", "2009-05-18"),
]


def write_meters(path):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Meters"
    ws["A1"] = "Town of Quillfield Water Department: meter register"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Revenue and process meters on the SCADA system and in the billing system. Updated September 2026."
    ws["A2"].font = Font(italic=True, size=9, color="555555")
    head = ["Meter", "Type", "Location", "Measures", "Units", "In service"]
    for j, h in enumerate(head, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE7EE")
    for i, r in enumerate(METERS, 5):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEF", (10, 26, 44, 40, 7, 11)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    tags = wb.create_sheet("SCADA tags")
    tags["A1"] = "Daily totals exported from the plant historian"
    tags["A1"].font = Font(bold=True)
    for j, h in enumerate(["Tag", "Meter", "Description"], 1):
        c = tags.cell(row=3, column=j, value=h)
        c.font = Font(bold=True)
    for i, (t, m, d) in enumerate([("WTP_RAW_TOTAL", "WTP-RAW", "Raw water from river intake, daily total"),
                                   ("WTP_FIN_TOTAL", "WTP-FIN", "Finished water, clearwell effluent, daily total"),
                                   ("HR_WELLS_TOTAL", "HR-WF", "Hollow Road wellfield to system, daily total (radio "
                                                               "telemetry)")], 4):
        for j, v in enumerate((t, m, d), 1):
            tags.cell(row=i, column=j, value=v)
    for col, w in zip("ABC", (18, 10, 60)):
        tags.column_dimensions[col].width = w
    wb.save(path)


# ---------------------------------------------------------------- the dictionary and the folder's contents
DICT = [
    ("gauge", "agency", "Agency operating the gage (SHS, State Hydrologic Survey)"),
    ("gauge", "station", "Station number"),
    ("gauge", "date", "Day, local standard time"),
    ("gauge", "discharge_cfs", "Daily mean discharge, cubic feet per second"),
    ("gauge", "qual", "Qualification code: A approved, P provisional; :e appended (A:e, P:e) when the value was "
                      "estimated"),
    ("notices", "notice_no", "Select Board drought notice number"),
    ("notices", "posted", "Date the notice was posted"),
    ("notices", "stage", "Drought stage under the drought management plan"),
    ("notices", "action", "Declared or Rescinded"),
    ("notices", "in_effect_from", "First day a declared stage is in effect"),
    ("notices", "in_effect_through", "Last day a rescinded stage is in effect"),
    ("notices", "basis", "Reading the notice relied on"),
    ("scada", "RecordID", "Historian record number"),
    ("scada", "Date", "Day of the daily total, midnight to midnight"),
    ("scada", "Tag", "Historian tag (see the meter register)"),
    ("scada", "Description", "Tag description"),
    ("scada", "DailyTotal", "Total for the day"),
    ("scada", "Units", "MG, million gallons"),
    ("elmhollow", "read_date", "Day the consumption is for"),
    ("elmhollow", "meter_id", "Meter (see the meter register)"),
    ("elmhollow", "consumption_kgal", "Consumption for the day, thousands of gallons"),
    ("elmhollow", "read_type", "Actual, or Estimated when the radio read failed"),
    ("welllog", "log_date", "Day of the operator's round"),
    ("welllog", "reading_time", "Time the totalizer was read"),
    ("welllog", "totalizer_gal", "Hollow Road wellfield totalizer, gallons"),
    ("welllog", "remarks", "Operator's remarks"),
    ("precip", "station", "Cooperative weather station"),
    ("precip", "year", "Year"),
    ("precip", "month", "Month"),
    ("precip", "precip_in", "Monthly precipitation, inches"),
    ("precip", "attributes", "M: missing (more than nine days missing in the month)"),
]

INDEX = [
    ("gauge", W.SURVEY, "Daily mean discharge, Varrow River at Quillfield, April 1951 to October 2026", "2026-10-16"),
    ("station", W.SURVEY, "Station description for the gauge", "2025-11-20"),
    ("reports", W.SOCIETY, "Water Works annual reports 1894 to 1950: pumping stopped under the mill agreement",
     "2019-06-03"),
    ("guidance", W.BOARD, "Guidance WR-14, design drought flows for small public water systems", "2023-03-15"),
    ("regression", W.BOARD, "Regional low-flow regression equations, Technical Report 2019-3", "2019-09-30"),
    ("permit", W.BOARD, "Water withdrawal permit WD-0417", "2021-03-03"),
    ("plan", "Quillfield Water Department", "Drought management plan, 2024 revision", "2024-05-14"),
    ("offer", "City of Castlewick", "Offer of firm interconnect capacity, 2027 to 2031", "2026-10-06"),
    ("notices", "Quillfield Select Board", "Drought stage notices since 2012", "2026-10-15"),
    ("scada", "Quillfield Water Department", "Mill Street plant historian, daily totals since 2012", "2026-10-16"),
    ("elmhollow", "Quillfield Water Department", "Elm Hollow master meter, daily consumption since 2012", "2026-10-16"),
    ("welllog", "Quillfield Water Department", "Hollow Road operator's round sheets since June 2023", "2026-10-16"),
    ("meters", "Quillfield Water Department", "Meter register and historian tags", "2026-09-28"),
    ("precip", "Quillfield 2 NW cooperative observer", "Monthly precipitation, 1895 to 2025", "2026-01-12"),
    ("dict", "Quillfield Water Department", "Field definitions for the data files in this folder", "2026-10-16"),
]


UNITS = {"discharge_cfs": "ft3/s", "DailyTotal": "MG", "consumption_kgal": "1,000 gal", "totalizer_gal": "gal",
         "precip_in": "in"}


def write_dictionary(path):
    rows = [[W.FILES[k], f, d, UNITS.get(f, "")] for k, f, d in DICT]
    return P.write_csv(path, ["file_name", "column_name", "description", "units"], rows)


def write_index(path, counts):
    rows = []
    for k, src, desc, d in INDEX:
        rows.append([W.FILES[k], src, desc, d, counts.get(k, "")])
    return P.write_csv(path, ["file", "from", "contents", "as_of", "rows"], rows)
