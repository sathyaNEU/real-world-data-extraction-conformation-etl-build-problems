"""Independent verifier for task116. Reads only the shipped files in target/ and the golden figures, on its own
code path (csv module, pyarrow, plain dictionaries), and recomputes every graded figure. Exit status 1 on any
disagreement."""
import csv
import datetime as dt
import json
import sys
from collections import defaultdict
from pathlib import Path

import pyarrow.parquet as pq
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
T = HERE / "target"
GOLD = json.loads((HERE / "generator" / "golden_figures.json").read_text())
AS_AT = dt.date(2026, 10, 1)
RATE = 4.05
fails = []


def same(a, b, what):
    if a != b:
        fails.append("%s: verifier %r, golden %r" % (what, a, b))


def rows(name, enc="utf-8"):
    with open(T / name, newline="", encoding=enc) as f:
        return list(csv.DictReader(f))


def d(s, fmt="%Y-%m-%d"):
    return dt.datetime.strptime(s, fmt).date() if s else None


est = rows("edubasealldata20261007.csv", "latin-1")
links = rows("links_edubasealldata20261007.csv", "latin-1")
groups = rows("allgroupsdata20261007.csv", "latin-1")
book = rows("tb_agreements_2026-10-01.csv")
crm = rows("tb_crm_opportunities_2022-2026.csv")
ledger = rows("tb_billing_ledger_2025-04_2026-09.csv")
remit = rows("orrin_remittance_2026Q3.csv")

E = {int(r["URN"]): r for r in est}
status = {u: r["EstablishmentStatus (name)"] for u, r in E.items()}
succ = defaultdict(list)
for r in links:
    if r["LinkType"] == "Successor":
        succ[int(r["URN"])].append(int(r["LinkURN"]))


def current(u):
    seen, cur = set(), u
    while status.get(cur) == "Closed" and len(succ.get(cur, [])) == 1 and cur not in seen:
        seen.add(cur)
        cur = succ[cur][0]
    return cur if status.get(cur) == "Open" else u


mat = {int(g["Group UID"]): g for g in groups if g["Group Type"] == "Multi-academy trust" and g["Group Status"] == "Open"}
open_in = defaultdict(list)
for u, r in E.items():
    if r["EstablishmentStatus (name)"] == "Open" and r["Trusts (code)"]:
        t = int(float(r["Trusts (code)"]))
        if t in mat:
            open_in[t].append(u)


def pupils(u):
    v = E[u]["NumberOfPupils"]
    return int(float(v)) if v else 0


wb = load_workbook(T / "trust_licence_pilot_2026.xlsx", read_only=True)
offers = [dict(zip([c.value for c in next(wb["Offers"].iter_rows(max_row=1))], [c.value for c in r]))
          for r in wb["Offers"].iter_rows(min_row=2)]
sched = [dict(zip(["trust", "urn", "academy", "pupils"], [c.value for c in r])) for r in wb["Schedules"].iter_rows(min_row=2)]
pilot_ch = set(o["Companies House no."] for o in offers)

# ------------------------------------------------------------------ the schools each agreement bills
# The export keys an agreement to the school whose account it sits on; the invoice for its current term names every
# school it bills, matched by name to the open academies of that school's trust.
name_in_trust = {}
for u, r in E.items():
    if r["EstablishmentStatus (name)"] == "Open" and r["Trusts (code)"]:
        name_in_trust[(int(float(r["Trusts (code)"])), r["EstablishmentName"])] = u
billed_names = {}
for r in ledger:
    if r["document_type"] == "Invoice" and " subscription " in r["description"] and ": " in r["description"]:
        opens = r["description"].split(" subscription ")[1].split(" to ")[0]
        billed_names[(r["agreement_ref"], d(opens, "%d/%m/%Y"))] = r["description"].split(": ", 1)[1].split("; ")


# every name on a group invoice must resolve to exactly one open academy of its lead school's trust
for (ref, _), names in billed_names.items():
    if len(names) > 1:
        lead_rows = [r for r in book if r["agreement_ref"] == ref and r["urn"]]
        tc = E[current(int(lead_rows[0]["urn"]))]["Trusts (code)"] if lead_rows else ""
        for n in names:
            hits = [u for u, r in E.items() if r["EstablishmentStatus (name)"] == "Open" and r["EstablishmentName"] == n
                    and r["Trusts (code)"] == tc]
            if len(hits) != 1:
                fails.append("group invoice %s names '%s', which is not one open academy of the lead's trust" % (ref, n))


def covers(r):
    lead = current(int(r["urn"]))
    out = [lead]
    names = billed_names.get((r["agreement_ref"], d(r["term_start"])), [])
    tc = E[lead]["Trusts (code)"]
    if len(names) > 1 and tc:
        for n in names:
            u = name_in_trust.get((int(float(tc)), n))
            if u is not None and u not in out:
                out.append(u)
    return out


# ------------------------------------------------------------------ agreements in force at 1 October
seen_ref, inforce = set(), []
for r in book:
    if r["status"] != "In force" or r["agreement_ref"] in seen_ref:
        continue
    if d(r["term_start"]) <= AS_AT <= d(r["term_end"]):
        seen_ref.add(r["agreement_ref"])
        inforce.append(r)

# ------------------------------------------------------------------ Tallybridge's own people sign nothing for a customer
import re
staff = set(re.findall(r"\*\*([A-Z][a-z]+ [A-Z][A-Za-z']+)\*\*", (T / "trust_round_thread.md").read_text()))
staff |= set(o["owner"] for o in crm)
if len(staff) < 5 or staff & set(r["signed_by"] for r in book):
    fails.append("a Tallybridge name signs a customer agreement: %s" % sorted(staff & set(r["signed_by"] for r in book)))

# ------------------------------------------------------------------ the main call
tl = [o for o in crm if o["opportunity_type"] == "Trust Licence" and o["stage"] in ("Closed Won", "Closed Lost")]
won, closed = sum(o["stage"] == "Closed Won" for o in tl), len(tl)
same([won, closed], GOLD["round"]["sign_rate"], "sign rate")
by_trust = defaultdict(list)
for r in inforce:
    if not r["urn"]:
        continue
    cur = current(int(r["urn"]))
    tc = E[cur]["Trusts (code)"]
    if E[cur]["EstablishmentStatus (name)"] == "Open" and tc and int(float(tc)) in mat:
        by_trust[int(float(tc))].append((cur, r))
cands = []
for uid, lst in by_trust.items():
    if mat[uid]["Companies House Number"].zfill(8) in pilot_ch or len(set(u for _, r in lst for u in covers(r))) < 2:
        continue
    lic = RATE * sum(pupils(u) for u in open_in[uid])
    net = lic - sum(float(r["annual_value_gbp"]) for _, r in lst)
    # one person is one CRM contact, however the name was typed
    signers = set(r["signed_by_contact_id"] for _, r in lst)
    cands.append((round(net * won / closed), mat[uid]["Group Name"], net, len(signers) == 1))
central = sorted([c for c in cands if c[3]], key=lambda c: (-c[0], c[1]))
same([c[1] for c in central[:10]], GOLD["round"]["slate"], "slate")
same([c[0] for c in central[:10]], GOLD["round"]["planned"], "planned values")
same(sum(c[0] for c in central[:10]), GOLD["round"]["total"], "round total")
same(central[10][1], GOLD["round"]["first_below"], "first below the line")
same(central[9][0] - central[10][0], GOLD["round"]["gap"], "gap at the cut")
same([c[1] for c in central[10:15]], [b["trust"] for b in GOLD["round"]["behind"]], "five behind")
run, acc = [], 0
for c in central[:10]:
    acc += c[0]
    run.append(acc)
same(run, GOLD["round"]["running"], "running totals")

# ------------------------------------------------------------------ the supporting claims in the write-up
devolved = sorted([c for c in cands if not c[3]], key=lambda c: (-c[0], c[1]))
same(dict(trusts=len(cands), central=len(central), devolved=len(devolved)), GOLD["round"]["eligible"], "eligible trusts")
spelled = mixed = 0
for uid, lst in by_trust.items():
    acad = set(u for _, r in lst for u in covers(r))
    people, names = set(r["signed_by_contact_id"] for _, r in lst), set(r["signed_by"] for _, r in lst)
    if len(acad) >= 2 and len(people) == 1 and len(names) > 1:
        spelled += 1
    if len(acad) >= 2 and 1 < len(people) < len(acad):
        mixed += 1
if not (spelled >= 3 and mixed >= 3):
    fails.append("expected trusts with one contact under two spellings and trusts with one contact for some academies")
same([dict(trust=c[1], planned_if_signed=c[0]) for c in devolved[:2]], GOLD["round"]["rivals"], "rivals named")
if not all(c[0] > central[9][0] for c in devolved[:2]):
    fails.append("a named rival has less room to grow than the tenth trust")
uid_of_ch = {g["Companies House Number"].zfill(8): u for u, g in mat.items()}


def one_signer(uid, rows_):
    lst = [r for r in rows_ if r["urn"] and E[current(int(r["urn"]))]["Trusts (code)"]
           and int(float(E[current(int(r["urn"]))]["Trusts (code)"])) == uid]
    return len(set(current(int(r["urn"])) for r in lst)), len(set(r["signed_by_contact_id"] for r in lst))


cl = defaultdict(lambda: [0, 0])
for o in crm:
    if o["opportunity_type"] != "Cluster Licence":
        continue
    n_acad, n_sign = one_signer(uid_of_ch[o["companies_house_no"].zfill(8)], inforce)
    if n_acad < 2 or n_sign not in (1, n_acad):
        fails.append("cluster trust %s not classifiable" % o["account_name"])
    k = "central" if n_sign == 1 else "devolved"
    cl[k][0] += 1
    cl[k][1] += o["stage"] == "Closed Won"
same({k: dict(offered=v[0], won=v[1]) for k, v in cl.items()}, GOLD["round"]["cluster"], "cluster record")
reproduced = 0
for o in offers:
    uid = uid_of_ch[o["Companies House no."]]
    rep = [r for r in (book if o["Outcome"] == "Signed" else inforce)
           if r["urn"] and (r["status"] == "Replaced by Trust Licence" or o["Outcome"] != "Signed")
           and E[current(int(r["urn"]))]["Trusts (code)"] and int(float(E[current(int(r["urn"]))]["Trusts (code)"])) == uid]
    lic = RATE * sum(pupils(u) for u in open_in[uid])
    val = sum(float(r["annual_value_gbp"]) for r in rep)
    got = (len(open_in[uid]), sum(pupils(u) for u in open_in[uid]), round(lic, 2), len(rep), round(val, 2), round(lic - val, 2))
    want = (o["Academies priced"], o["Pupils priced"], round(o["Licence value (GBP)"], 2), o["Agreements replaced"],
            round(o["Agreements replaced: annual value (GBP)"], 2), round(o["Net new recurring revenue (GBP)"], 2))
    reproduced += got == want
    if len(set(r["signed_by_contact_id"] for r in rep)) != 1 or len(set(r["signed_by"] for r in rep)) != 1:
        fails.append("pilot trust %s is not one contact under one name" % o["Trust"])
same(reproduced, len(offers), "pilot offers reproduced")
recorded = 0
for o in offers:
    uid = uid_of_ch[o["Companies House no."]]
    rep = [r for r in (book if o["Outcome"] == "Signed" else inforce)
           if r["urn"] and (r["status"] == "Replaced by Trust Licence" or o["Outcome"] != "Signed")
           and E[int(r["urn"])]["Trusts (code)"] and int(float(E[int(r["urn"])]["Trusts (code)"])) == uid]
    val = sum(float(r["annual_value_gbp"]) for r in rep)
    recorded += round(RATE * sum(pupils(u) for u in open_in[uid]) - val, 2) == round(o["Net new recurring revenue (GBP)"], 2)
if recorded == len(offers):
    fails.append("recorded URNs reproduce the pilot too")
same(len(offers), 12, "pilot offers")

# ------------------------------------------------------------------ ask A: schools on Tallybridge at 1 October
def band(u):
    return "Primary" if E[u]["PhaseOfEducation (name)"] in ("Primary", "Middle deemed primary") else "Secondary"


def ex_vat(r):
    """The old system (L numbers) holds amounts with VAT in them and the VAT within them in vat_gbp; the new platform
    holds them net."""
    a = float(r["amount_gbp"])
    return round(a - float(r["vat_gbp"]), 2) if r["document_no"].startswith("L") else a


PRICE = {("2025-26", "Primary"): 4.40, ("2025-26", "Secondary"): 5.10,
         ("2026-27", "Primary"): 4.60, ("2026-27", "Secondary"): 5.30}
term_bill = {}
for r in ledger:
    if r["document_type"] == "Invoice" and " subscription " in r["description"]:
        opens = r["description"].split(" subscription ")[1].split(" to ")[0]
        term_bill[(r["agreement_ref"], d(opens, "%d/%m/%Y"))] = ex_vat(r)
stale = 0
school_rows = []
for r in inforce:
    if r["billing_basis"] == "Per pupil":
        cur = current(int(r["urn"]))
        p_, v_ = int(r["pupils_billed"]), float(r["annual_value_gbp"])
        billed = term_bill[(r["agreement_ref"], d(r["term_start"]))]
        if abs(billed - v_) > 0.005:
            # the CRM still carries the term before; the invoice for the current term is what it bills
            yr = "2026-27" if d(r["term_start"]) >= dt.date(2026, 4, 1) else "2025-26"
            n = billed / PRICE[(yr, r["price_band"])]
            if abs(n - round(n)) > 1e-6:
                fails.append("%s: current invoice is not a whole number of pupils at the price list" % r["agreement_ref"])
            p_, v_, stale = int(round(n)), billed, stale + 1
        school_rows.append((cur, E[cur]["EstablishmentTypeGroup (name)"], r["price_band"], p_, v_))
        for u in covers(r)[1:]:
            school_rows.append((u, E[u]["EstablishmentTypeGroup (name)"], r["price_band"], 0, 0.0))
    elif r["billing_basis"] == "Trust Licence":
        for s in sched:
            if s["trust"] == r["account_name"]:
                u = int(s["urn"])
                school_rows.append((u, E[u]["EstablishmentTypeGroup (name)"], band(u), int(s["pupils"]),
                                    RATE * int(s["pupils"])))
dfe = {}
for u, r in E.items():
    if r["EstablishmentNumber"]:
        dfe.setdefault("%d%04d" % (int(r["LA (code)"]), int(float(r["EstablishmentNumber"]))), []).append(u)
    if r["PreviousLA (code)"] and r["PreviousEstablishmentNumber"]:
        dfe.setdefault("%d%04d" % (int(float(r["PreviousLA (code)"])), int(float(r["PreviousEstablishmentNumber"]))), []).append(u)
fed_seen = set()
for r in remit:
    if r["reseller_ref"] == "TOTAL":
        continue
    if not (d(r["contract_start"], "%d/%m/%Y") <= AS_AT <= d(r["contract_end"], "%d/%m/%Y")):
        continue
    keep = 1 - float(r["reseller_margin_pct"]) / 100
    if not r["dfe_number"]:
        for s in sched:
            if s["trust"] == r["establishment_name"]:
                u = int(s["urn"])
                school_rows.append((u, E[u]["EstablishmentTypeGroup (name)"], band(u), int(s["pupils"]),
                                    RATE * int(s["pupils"]) * keep))
        continue
    u = [x for x in dfe[r["dfe_number"]] if E[x]["EstablishmentStatus (name)"] == "Open"][0]
    # a contract covering several schools is listed against each of them: its pupils and value count once
    first = r["contract_ref"] not in fed_seen
    fed_seen.add(r["contract_ref"])
    school_rows.append((u, E[u]["EstablishmentTypeGroup (name)"], band(u), int(r["pupils"]) if first else 0,
                        float(r["annual_contract_value_gbp"]) * keep if first else 0.0))
cell = defaultdict(lambda: [set(), 0, 0.0])
for u, t, b, p, v in school_rows:
    c = cell["%s|%s" % (t, b)]
    c[0].add(u)
    c[1] += p
    c[2] += v
for k, g in GOLD["book"].items():
    c = cell[k]
    same(dict(pupils=c[1], revenue=int(round(c[2]))), g, "book cell " + k)

# ------------------------------------------------------------------ ask B: the pilot licences
for o in offers:
    if o["Outcome"] != "Signed":
        continue
    trust = o["Trust"]
    lic = [r for r in book if r["billing_basis"] == "Trust Licence" and r["account_name"] == trust]
    if lic:
        fee = sum(ex_vat(r) for r in ledger if r["agreement_ref"] == lic[0]["agreement_ref"])
    else:
        fee = float([r for r in remit if r["establishment_name"] == trust][0]["annual_contract_value_gbp"])
    uid = [u for u, g in mat.items() if g["Group Name"] == trust][0]
    starts = {r["agreement_ref"]: r["term_start"] for r in book if r["status"] == "Replaced by Trust Licence" and r["urn"]
              and E[current(int(r["urn"]))]["Trusts (code)"]
              and int(float(E[current(int(r["urn"]))]["Trusts (code)"])) == uid}
    # the invoice for the term each agreement was in when the licence replaced it: its description opens on the
    # agreement's term start; old-system credit notes against it carry only the invoice number
    cut = set()
    for r in ledger:
        if r["document_type"] == "Invoice" and r["agreement_ref"] in starts:
            opens = r["description"].split(" subscription ")[1].split(" to ")[0]
            if d(opens, "%d/%m/%Y") == d(starts[r["agreement_ref"]]):
                cut.add(r["document_no"])
    credit = -sum(ex_vat(r) for r in ledger if r["document_type"] == "Credit note" and r["reference"] in cut
                  and d(r["document_date"]) >= dt.date(2026, 9, 1))
    same(dict(fee=round(fee, 2), credit=round(credit, 2), net=round(fee - credit, 2)), GOLD["pilot"][trust], "pilot " + trust)

# ------------------------------------------------------------------ ask C: usage per hundred billed
u = pq.read_table(T / "tb_weekly_usage_2025-09_2026-10.parquet").to_pylist()
u.sort(key=lambda r: (r["loaded_at"], r["school_ref"]), reverse=True)   # the latest load of a school's week stands
week0 = dt.date(2026, 9, 7)
weeks = set(week0 + dt.timedelta(days=7 * i) for i in range(6))
seen = set()
acts = defaultdict(int)
place = {}
for urn, t, b, p, v in school_rows:
    place[urn] = "%s|%s" % (t, b)
for r in u:
    # each weekly load reports one Monday-to-Sunday week, named by year and ISO week in its batch id
    _, yy, ww = r["batch_id"].split("-")[:3]
    monday = dt.date.fromisocalendar(int(yy), int(ww), 1)
    key = (r["school_ref"], monday)
    if key in seen:
        continue
    seen.add(key)
    if monday not in weeks:
        continue
    cur = current(int(r["urn"]))
    if cur in place:
        acts[place[cur]] += r["active_pupils"]
for k, g in GOLD["usage"].items():
    p = cell[k][1]
    same(round(100.0 * acts[k] / 6.0 / p, 1) if p else None, g, "usage " + k)

if fails:
    print("VERIFY FAILED (%d):" % len(fails))
    for f in fails:
        print("  " + f)
    sys.exit(1)
print("VERIFY PASSED: round %s (%d trusts) at %d of %d, %d eligible (%d central), cluster %s, %d of %d pilot offers "
      "reproduced (%d on recorded URNs), %d book cells (%d agreements billed off the invoice), %d pilot licences, "
      "%d usage cells"
      % (GOLD["round"]["total"], len(GOLD["round"]["slate"]), won, closed, len(cands), len(central), dict(cl), reproduced,
         len(offers), recorded, len(GOLD["book"]), stale, len(GOLD["pilot"]), len(GOLD["usage"])))
