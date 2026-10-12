"""Write the evidence pack into target/."""
import datetime as dt
import shutil
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

import docs
import extracts
import frames
from world import CACHE, REG_FILES, rng

D = dt.date
FILES = {
    "book": "tb_agreements_2026-10-01.csv",
    "crm": "tb_crm_opportunities_2022-2026.csv",
    "ledger": "tb_billing_ledger_2025-04_2026-09.csv",
    "usage": "tb_weekly_usage_2025-09_2026-10.parquet",
    "pilot": "trust_licence_pilot_2026.xlsx",
    "remit": "orrin_remittance_2026Q3.csv",
    "terms": "tallybridge_trust_licence_terms_2026.pdf",
    "prices": "tallybridge_price_list_2026-27.pdf",
    "plan": "trust_partnerships_plan_2027.docx",
    "notes": "revenue_ops_export_guide.md",
    "thread": "trust_round_thread.md",
    "finance": "finance_note_billing_cutover.txt",
    "sources": "about_this_folder.txt",
}


def pilot_workbook(path, reg, A):
    wb = Workbook()
    ws = wb.active
    ws.title = "Offers"
    head = ["Trust", "Companies House no.", "Request received", "Priced on", "Academies priced", "Pupils priced",
            "Licence value (GBP)", "Agreements replaced", "Agreements replaced: annual value (GBP)",
            "Net new recurring revenue (GBP)", "Outcome", "Decided on", "Licence start"]
    ws.append(head)
    order = sorted(A.licences, key=lambda L: rng("pilotorder", L["uid"]).random())
    for L in order:
        uid = L["uid"]
        r = rng("pilotdates", uid)
        req = D(2026, 2, 2) + dt.timedelta(days=int(r.integers(0, 80)))
        dec = D(2026, 6, 22) + dt.timedelta(days=int(r.integers(0, 50)))
        ws.append([reg.group_name[uid], reg.group_ch.get(uid), req, D(2026, 6, 15), L["academies"], L["pupils"],
                   L["licence_value"], L["replaced_n"], L["replaced_value"], L["net"], L["outcome"], dec,
                   D(2026, 9, 1) if L["outcome"] == "Signed" else None])
    for row in ws.iter_rows(min_row=2):
        for c in (row[2], row[3], row[11], row[12]):
            c.number_format = "dd/mm/yyyy"
        for c in (row[6], row[8], row[9]):
            c.number_format = "#,##0.00"
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [44, 14, 12, 11, 10, 10, 13, 11, 16, 15, 10, 11, 11]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    s2 = wb.create_sheet("Schedules")
    s2.append(["Trust", "URN", "Academy", "Pupils priced"])
    for L in A.licences:
        if L["outcome"] != "Signed":
            continue
        for urn, pup in zip(L["academy_urns"], L["academy_pupils"]):
            s2.append([reg.group_name[L["uid"]], urn, reg.info[urn]["EstablishmentName"], pup])
    for c in s2[1]:
        c.font = Font(bold=True)
    s2.column_dimensions["A"].width = 44
    s2.column_dimensions["C"].width = 48
    s3 = wb.create_sheet("Notes")
    for line in ["Offers made to the trusts that asked for trust pricing, spring 2026.",
                 "All twelve priced on 15 June 2026 on pupil numbers as GIAS published them that day.",
                 "Net new recurring revenue is the licence value less the annual value of the academy agreements the "
                 "licence replaces.",
                 "Schedules list the academies covered by each licence signed, as priced.",
                 "Owner: trust partnerships (Pamela Andrews)."]:
        s3.append([line])
    s3.column_dimensions["A"].width = 110
    wb.save(path)


def write(T, reg, A):
    T.mkdir(parents=True, exist_ok=True)
    for f in T.iterdir():
        if f.is_file():
            f.unlink()
    for name in REG_FILES.values():
        if "membership" in name:
            continue
        shutil.copyfile(CACHE / name, T / name)
    book = frames.book_frame(A.rows)
    book.to_csv(T / FILES["book"], index=False)
    crm = extracts.crm(reg, A)
    crm.to_csv(T / FILES["crm"], index=False)
    led = extracts.ledger(reg, A)
    led.to_csv(T / FILES["ledger"], index=False)
    rem = extracts.remittance(reg, A)
    rem.to_csv(T / FILES["remit"], index=False)
    us = extracts.usage(reg, A)
    us.to_parquet(T / FILES["usage"], index=False)
    pilot_workbook(T / FILES["pilot"], reg, A)
    docs.licence_terms(T / FILES["terms"])
    docs.price_list(T / FILES["prices"])
    docs.plan_doc(T / FILES["plan"])
    docs.data_notes(T / FILES["notes"])
    docs.thread(T / FILES["thread"])
    docs.finance_note(T / FILES["finance"])
    counts = dict(est=len(reg.est), links=len(reg.links), groups=len(reg.groups), book=len(book), crm=len(crm),
                  ledger=len(led), usage=len(us), remit=len(rem) - 1)
    docs.sources(T / FILES["sources"], counts)
    return dict(book=book, crm=crm, ledger=led, remit=rem, usage=us)
