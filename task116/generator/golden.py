"""The golden deliverables, written from the verified figures and the shipped files only."""
import datetime as dt
import io
import json
import re
import zipfile
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

import golden_calc as G

HERE = Path(__file__).resolve().parent
TASK = HERE.parent
OUT = TASK / "golden"
AUTHOR = "Pamela Andrews"
CREATED, SAVED = dt.datetime(2026, 10, 19, 9, 12, 0), dt.datetime(2026, 10, 21, 16, 40, 0)
BLUE, GREY, INK, INK2 = "#2a78d6", "#b9b7b0", "#0b0b0b", "#52514e"
TYPE_LABEL = {"Academies": "Academies", "Free Schools": "Free schools",
              "Local authority maintained schools": "Local authority maintained schools",
              "Special schools": "Special schools"}


def title_case(name):
    """Register names are upper case; set them in title case, keeping initialisms."""
    small = {"OF", "AND", "THE", "FOR", "IN"}
    out = []
    for i, w in enumerate(name.split()):
        core = w.strip("().,&'")
        if w == "LTD":
            out.append("Ltd")
        elif w == "ST":
            out.append("St")
        elif core and len(core) <= 4 and not any(v in core for v in "AEIOUY"):
            out.append(w)
        elif "." in core:
            out.append(w)
        elif i and w in small:
            out.append(w.lower())
        else:
            out.append("(" + w[1:].capitalize() if w.startswith("(") else w.capitalize())
    return " ".join(out)


def money(x):
    return int(round(x))


def build(T=TASK / "target"):
    f = G.write_figures(T, HERE / "golden_figures.json")
    d, m, sch, A, B, C = G.compute(T)
    OUT.mkdir(exist_ok=True)
    workbook(f, m, OUT / "trust_round_spring2027.xlsx")
    deck(f, OUT / "trust_round_spring2027.pptx")
    return f


# ---------------------------------------------------------------- workbook
HEAD = Font(bold=True)
THIN = Side(style="thin", color="999999")


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _header(ws, row, labels):
    for j, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=j, value=lab)
        c.font = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
        c.border = Border(bottom=THIN)


def workbook(f, m, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Spring round"
    r = f["round"]
    ws["A1"] = "Trust Licence: spring 2027 round"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Ten trusts, in the order Leah and Joel take the licence to them, for licences from 1 September "
                "2027. Pupils as GIAS published them on 7 October 2026.")
    cols = ["Priority", "Trust", "Companies House no.", "Open academies", "Academies on an agreement",
            "Pupils on roll", "Licence fee (£)", "Agreements it replaces (£)", "New recurring revenue (£)",
            "Planned value (£)", "Running total (£)"]
    _header(ws, 4, cols)
    sl = m["slate"]
    for i, (_, x) in enumerate(sl.iterrows(), 1):
        ws.append([i, title_case(x["trust"]), x["companies_house_no"], int(x["academies"]), int(x["subscribing"]),
                   int(x["pupils"]), money(x["licence_value"]), money(x["replaced"]),
                   money(x["new_recurring_revenue"]), int(x["planned_if_signed"]), int(r["running"][i - 1])])
    tr = ws.max_row + 1
    ws.cell(row=tr, column=2, value="Round on the 2027-28 plan").font = HEAD
    c = ws.cell(row=tr, column=10, value=int(r["total"]))
    c.font = HEAD
    c.border = Border(top=THIN)
    ws.cell(row=tr + 2, column=1, value="Below the line").font = HEAD
    b0 = tr + 3
    _header(ws, b0, cols[:10])
    for i, (_, x) in enumerate(m["central"].iloc[10:15].iterrows(), 11):
        ws.append([i, title_case(x["trust"]), x["companies_house_no"], int(x["academies"]), int(x["subscribing"]),
                   int(x["pupils"]), money(x["licence_value"]), money(x["replaced"]),
                   money(x["new_recurring_revenue"]), int(x["planned_if_signed"])])
    gap_row = ws.max_row + 1
    ws.cell(row=gap_row, column=2, value="Gap from the tenth trust to the first below the line")
    ws.cell(row=gap_row, column=10, value=int(r["gap"]))
    # trusts with more room that a trust-wide licence cannot reach
    s0 = gap_row + 2
    ws.cell(row=s0, column=1, value=("Not eligible under the round plan: more room than the tenth trust, but its "
                                     "academies' agreements were signed by different people")).font = HEAD
    _header(ws, s0 + 1, ["", "Trust", "Companies House no.", "Open academies", "Academies on an agreement",
                         "People who signed them", "Licence fee (£)", "Agreements it replaces (£)",
                         "New recurring revenue (£)", "Planned value (£)"])
    t = m["table"]
    aside = t[(t["mode"] == "devolved") & (t["planned_if_signed"] > r["planned"][-1])]
    for _, x in aside.iterrows():
        ws.append(["", title_case(x["trust"]), x["companies_house_no"], int(x["academies"]), int(x["subscribing"]),
                   int(x["signatories"]), money(x["licence_value"]), money(x["replaced"]),
                   money(x["new_recurring_revenue"]), "nil"])
    cl = r["cluster"]
    sr = r["sign_rate"]
    notes = [
        "Licence fee: £4.05 a pupil on every academy GIAS shows open in the trust (Trust Licence terms, 3).",
        "New recurring revenue: the licence fee less the annual value of the agreements it replaces (terms, 4).",
        "Planned value: new recurring revenue x %d/%d, the sign rate of the Trust Licence offers closed to date "
        "(%d signed of %d: the twelve pilot offers and the offer declined on 30 September), rounded to the pound. "
        "The round is the sum of the ten." % (sr[0], sr[1], sr[0], sr[1]),
        "Eligible (round plan): an open multi-academy trust with at least two academies on an agreement in force and "
        "the same person having signed every one; pilot trusts are out. Of the %d trusts outside the pilot with two or "
        "more academies on an agreement, %d qualify. For reference, Cluster Licence offers 2023-24: %d of %d trusts of that kind signed; "
        "%d of %d trusts whose academies each signed their own agreement did."
        % (r["eligible"]["trusts"], r["eligible"]["central"], cl["central"]["won"], cl["central"]["offered"],
           cl["devolved"]["won"], cl["devolved"]["offered"]),
        "Academies converted after they first subscribed are matched to their academy record through GIAS links.",
        "Academies on an agreement: every school the invoice for the agreement's current term bills. Four agreements "
        "each bill three or four academies of one trust and sit in the CRM on the lead school's account; Bishop "
        "Hogarth and Derby Diocesan are two of them.",
    ]
    ws.cell(row=ws.max_row + 2, column=1, value="Definitions and sources are on the Notes sheet.").font = Font(italic=True)
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
        for c in row[6:11]:
            if isinstance(c.value, (int, float)):
                c.number_format = "#,##0"
        if isinstance(row[5].value, int) and row[5].value > 999:
            row[5].number_format = "#,##0"
    _widths(ws, [8, 46, 13, 10, 12, 10, 12, 14, 14, 12, 12])
    ws.freeze_panes = "C5"
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1

    # the book at 1 October
    s2 = wb.create_sheet("Book at 1 October")
    s2["A1"] = "Schools on Tallybridge at 1 October 2026"
    s2["A1"].font = Font(bold=True, size=13)
    s2["A2"] = ("Agreements in force, billed as the invoice for their current term bills them; academies on a Trust "
                "Licence; contracts held through the framework reseller (recurring revenue net of its 12% margin). "
                "Type as GIAS groups the school.")
    _header(s2, 4, ["Type of school", "Price band", "Pupils billed", "Recurring revenue (£)"])
    tot = [0, 0]
    for k, v in f["book"].items():
        typ, band = k.split("|")
        s2.append([TYPE_LABEL[typ], band, v["pupils"] if v["pupils"] else "none",
                   v["revenue"] if v["pupils"] else "none"])
        tot = [tot[0] + v["pupils"], tot[1] + v["revenue"]]
    s2.append(["All schools", "", tot[0], tot[1]])
    for c in s2[s2.max_row]:
        c.font = HEAD
        c.border = Border(top=THIN)
    for row in s2.iter_rows(min_row=5, max_row=s2.max_row):
        for c in row[2:4]:
            if isinstance(c.value, int):
                c.number_format = "#,##0"
            else:
                c.alignment = Alignment(horizontal="right")
    _widths(s2, [34, 12, 14, 18])
    s2.freeze_panes = "A5"

    # the pilot licences
    s3 = wb.create_sheet("Pilot licences")
    s3["A1"] = "Trust Licences signed in the 2026 pilot: billed for 2026-27"
    s3["A1"].font = Font(bold=True, size=13)
    s3["A2"] = ("Excluding VAT. Old-system documents (L numbers) hold their amount with VAT in it, the VAT field "
                "showing the VAT within it; new-platform documents hold the amount net. Credit: the credit notes raised "
                "from 1 September against the term each replaced agreement was in. The reseller invoices the licence "
                "bought through the framework.")
    _header(s3, 4, ["Trust", "Credited back when it replaced the schools' subscriptions (£)",
                    "Billed for 2026-27, net of that credit (£)"])
    for trust in sorted(f["pilot"]):
        v = f["pilot"][trust]
        s3.append([title_case(trust), money(v["credit"]), money(v["net"])])
    for row in s3.iter_rows(min_row=5, max_row=s3.max_row):
        for c in row[1:3]:
            c.number_format = "#,##0"
    _widths(s3, [50, 22, 22])
    s3.freeze_panes = "A5"

    # usage
    s4 = wb.create_sheet("Usage autumn 2026")
    s4["A1"] = "Weekly active pupils per hundred billed, first six weeks of the autumn term"
    s4["A1"].font = Font(bold=True, size=13)
    s4["A2"] = "Weeks beginning 7 September to 12 October 2026. Same schools as the book at 1 October."
    _header(s4, 4, ["Type of school", "Price band", "Weekly active pupils per hundred billed"])
    for k, v in f["usage"].items():
        typ, band = k.split("|")
        s4.append([TYPE_LABEL[typ], band, float(v) if v is not None else "none"])
    for row in s4.iter_rows(min_row=5, max_row=s4.max_row):
        if isinstance(row[2].value, float):
            row[2].number_format = "0.0"
    _widths(s4, [34, 12, 22])
    s4.freeze_panes = "A5"
    s5 = wb.create_sheet("Notes")
    s5["A1"] = "Notes"
    s5["A1"].font = Font(bold=True, size=13)
    lines = notes + [
        "Philip's view was to go where we have most room to grow. The trusts with the most room are listed as not "
        "eligible on the Spring round sheet: different people signed their academies' agreements.",
        "Schools on Tallybridge: agreements in force at 1 October 2026 (schools the register shows proposed to close "
        "included), academies on a Trust Licence, and contracts held through the framework reseller, which are "
        "counted at the annual contract value less its 12% margin.",
        "Agreements that renewed in April 2026 still carry the term before's pupils and value in the CRM; they are "
        "taken as the invoice for the current term bills them (%d agreements)." % f["book_invoiced"],
        "A federation buying through the framework holds one contract for its schools, and the reseller lists it "
        "against each school it supplies under one contract reference: its pupils and value count once.",
        "Weekly active pupils per hundred billed: the weekly active pupils over the six weeks, divided by six, per "
        "hundred pupils billed across every school in the book. Weeks are read from the load that reports them "
        "(batch_id); the trust reporting service labels its weeks by the Sunday before. Where a week was loaded "
        "again, the later load stands. A school with no row in a week had no pupil answering that week.",
        "Money is in whole pounds excluding VAT, carried in pence and rounded once.",
        "Sources: GIAS public extract of 7 October 2026; tb_agreements_2026-10-01.csv; trust_licence_pilot_2026.xlsx; "
        "tb_crm_opportunities_2022-2026.csv; tb_billing_ledger_2025-04_2026-09.csv; orrin_remittance_2026Q3.csv; "
        "tb_weekly_usage_2025-09_2026-10.parquet.",
        "Prepared by Pamela Andrews for the planning meeting on 12 November 2026.",
    ]
    for i, nline in enumerate(lines, 3):
        c = s5.cell(row=i, column=1, value=nline)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    _widths(s5, [120])
    wb.properties.creator = wb.properties.lastModifiedBy = AUTHOR
    wb.save(path)
    _repack(path, {})


# ---------------------------------------------------------------- deck
def chart_png(f):
    """The round built up in priority order: each trust a step from the running total before it to the running total
    after it, the first trust below the line set after the cut."""
    r = f["round"]
    names = [title_case(n) for n in r["slate"]]
    run = r["running"]
    starts = [0] + run[:-1]
    first = r["behind"][0]
    fig, ax = plt.subplots(figsize=(12.0, 6.6), dpi=150)
    y = list(range(12))[::-1]
    yt, yb = y[:10], y[11]
    ax.barh(yt, r["planned"], left=starts, color=BLUE, height=0.66, edgecolor="white", linewidth=1.2)
    for i, (yi, s0, v) in enumerate(zip(yt, starts, run)):
        ax.text(v + 1400, yi, "£{:,}".format(v), va="center", fontsize=9.5, color=INK)
        if i:
            ax.plot([s0, s0], [yi + 0.33, yi + 0.67], color="#9a988f", linewidth=0.8)
    total = r["total"]
    ax.plot([total, total], [yb - 0.45, yt[-1] + 0.33], color=INK2, linestyle=(0, (4, 3)), linewidth=1.1)
    ax.axhline(y[10], color="#c8c6bf", linewidth=0.9)
    ax.text(1200, y[10] + 0.08, "The cut: ten trusts, £{:,}".format(total), fontsize=8.5, color=INK2, va="bottom")
    ax.barh([yb], [first["planned"]], left=[total], color=GREY, height=0.66, edgecolor="white")
    ax.text(total + first["planned"] + 1400, yb,
            "£{:,}, £{:,} below the tenth".format(first["planned"], r["gap"]), va="center", fontsize=9.5, color=INK)
    ax.set_yticks(yt + [yb])
    ax.set_yticklabels(["%d  %s" % (i, n) for i, n in enumerate(names, 1)] + [title_case(first["trust"])],
                       fontsize=10, color=INK)
    ax.get_yticklabels()[-1].set_color(INK2)
    ax.set_xlim(0, (total + first["planned"]) * 1.38)
    ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, p: "£{:,.0f}".format(v)))
    ax.tick_params(axis="x", colors=INK2, labelsize=9)
    ax.tick_params(axis="y", length=0)
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.spines["bottom"].set_color("#c8c6bf")
    ax.grid(axis="x", color="#ecebe7", linewidth=0.8)
    ax.set_axisbelow(True)
    ax.set_xlabel("Planned value on the 2027-28 plan, running total in priority order (new recurring revenue x %d/%d)"
                  % tuple(r["sign_rate"]), fontsize=9.5, color=INK2)
    ax.set_title("Ten trusts build the round to £{:,} on the 2027-28 plan".format(total), loc="left", fontsize=15,
                 color=INK, pad=14, fontweight="bold")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor="white", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def deck(f, path):
    r = f["round"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank)
    tb = s1.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(1.0))
    p = tb.text_frame.paragraphs[0]
    p.text = "Spring round: ten trusts, £{:,} on the 2027-28 plan".format(r["total"])
    p.runs[0].font.size, p.runs[0].font.bold = Pt(28), True
    p.runs[0].font.color.rgb = RGBColor(0x0B, 0x0B, 0x0B)
    sub = s1.shapes.add_textbox(Inches(0.6), Inches(1.25), Inches(12), Inches(0.6)).text_frame.paragraphs[0]
    sub.text = ("Trust Licence approaches January to Easter, licences from 1 September 2027. Planned value is new "
                "recurring revenue at the sign rate of the offers closed to date, %d of %d." % tuple(r["sign_rate"]))
    sub.runs[0].font.size = Pt(13)
    sub.runs[0].font.color.rgb = RGBColor(0x52, 0x51, 0x4E)
    rows, cols = 11, 2
    tbl = s1.shapes.add_table(rows, cols, Inches(0.6), Inches(2.0), Inches(8.2), Inches(4.4)).table
    tbl.columns[0].width, tbl.columns[1].width = Inches(0.9), Inches(7.3)
    for j, h in enumerate(["Priority", "Trust"]):
        tbl.cell(0, j).text = h
    for i, n in enumerate(r["slate"], 1):
        tbl.cell(i, 0).text = str(i)
        tbl.cell(i, 1).text = title_case(n)
    for i in range(rows):
        for j in range(cols):
            for para in tbl.cell(i, j).text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(12)
                    run.font.bold = (i == 0)
    side = s1.shapes.add_textbox(Inches(9.2), Inches(2.0), Inches(3.6), Inches(4.4)).text_frame
    side.word_wrap = True
    lines = ["First below the line: %s, £%s, £%s short of the tenth."
             % (title_case(r["first_below"]), "{:,}".format(r["behind"][0]["planned"]), "{:,}".format(r["gap"])),
             "Not on the list: %s, %s and the other trusts with more room to grow. Different people signed their "
             "academies' agreements, and the round plan offers the licence only where one person signed them all."
             % (title_case(r["rivals"][0]["trust"]), title_case(r["rivals"][1]["trust"]))]
    for k, line in enumerate(lines):
        para = side.paragraphs[0] if k == 0 else side.add_paragraph()
        para.text = line
        para.runs[0].font.size = Pt(13)
        para.space_after = Pt(12)
    s2 = prs.slides.add_slide(blank)
    png = chart_png(f)
    pic = s2.shapes.add_picture(io.BytesIO(png), Inches(0.4), Inches(0.35), height=Inches(6.45))
    pic.left = int((prs.slide_width - pic.width) / 2)
    src = ("Source: GIAS extract of 7 October 2026; Tallybridge agreements at 1 October 2026; trust partnerships "
           "round plan, 8 October 2026; Trust Licence offers closed to date (%d signed of %d)." % tuple(r["sign_rate"]))
    for n, sl in enumerate(prs.slides, 1):
        foot = sl.shapes.add_textbox(Inches(0.6), Inches(7.0), Inches(11.4), Inches(0.35)).text_frame.paragraphs[0]
        foot.text = src
        foot.runs[0].font.size = Pt(9)
        foot.runs[0].font.color.rgb = RGBColor(0x6B, 0x69, 0x64)
        num = sl.shapes.add_textbox(Inches(12.3), Inches(7.0), Inches(0.5), Inches(0.35)).text_frame.paragraphs[0]
        num.text = str(n)
        num.runs[0].font.size = Pt(9)
        num.runs[0].font.color.rgb = RGBColor(0x6B, 0x69, 0x64)
    cp = prs.core_properties
    cp.author = cp.last_modified_by = AUTHOR
    cp.comments, cp.title, cp.revision = "", "", 1
    prs.save(path)
    _repack(path, _deck_app(prs))


def _deck_app(prs):
    """The template's extended properties describe an empty 4:3 deck; restate them for the deck as saved."""
    paras = [p for sl in prs.slides for sh in sl.shapes for tf in
             ([sh.text_frame] if sh.has_text_frame else
              [c.text_frame for row in sh.table.rows for c in row.cells] if sh.has_table else [])
             for p in tf.paragraphs if p.text.strip()]
    words = sum(len(p.text.split()) for p in paras)
    n = len(prs.slides)
    titles = "".join("<vt:lpstr>PowerPoint Presentation</vt:lpstr>" for _ in range(n))
    heads = ('<HeadingPairs><vt:vector size="4" baseType="variant"><vt:variant><vt:lpstr>Theme</vt:lpstr>'
             '</vt:variant><vt:variant><vt:i4>1</vt:i4></vt:variant><vt:variant><vt:lpstr>Slide Titles</vt:lpstr>'
             '</vt:variant><vt:variant><vt:i4>%d</vt:i4></vt:variant></vt:vector></HeadingPairs>' % n)
    parts = ('<TitlesOfParts><vt:vector size="%d" baseType="lpstr"><vt:lpstr>Office Theme</vt:lpstr>%s'
             '</vt:vector></TitlesOfParts>' % (n + 1, titles))
    sub = {tag: "<%s>%s</%s>" % (tag, val, tag) for tag, val in
           [("TotalTime", "52"), ("Words", str(words)), ("PresentationFormat", "Widescreen"),
            ("Paragraphs", str(len(paras))), ("Slides", str(n)), ("AppVersion", "16.0000")]}
    sub["HeadingPairs"], sub["TitlesOfParts"] = heads, parts
    return sub


def _repack(path, app):
    """Fix the save clock: core dates at the in-fiction times, every entry stamped at the save time."""
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        data = {i.filename: z.read(i.filename) for i in infos}
    core = data["docProps/core.xml"].decode()
    for tag, when in [("created", CREATED), ("modified", SAVED)]:
        core = re.sub(r"(<dcterms:%s[^>]*>)[^<]*(</dcterms:%s>)" % (tag, tag),
                      r"\g<1>%s\g<2>" % when.strftime("%Y-%m-%dT%H:%M:%SZ"), core)
    data["docProps/core.xml"] = core.encode()
    b = data["docProps/app.xml"].decode()
    for tag, repl in app.items():
        b = re.sub(r"<%s>.*?</%s>" % (tag, tag), repl, b, flags=re.S)
    data["docProps/app.xml"] = b.encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for i in infos:
            zi = zipfile.ZipInfo(i.filename, date_time=SAVED.timetuple()[:6])
            zi.compress_type, zi.external_attr = zipfile.ZIP_DEFLATED, i.external_attr
            z.writestr(zi, data[i.filename])


def scrub():
    import subprocess
    tool = TASK.parent / ".claude/skills/reduce-house-fixes/scripts/scrub_producer_metadata.py"
    subprocess.run(["python3", str(tool), str(OUT), "--apply", "--producer", AUTHOR, "--stamp", "2026-10-21",
                    "--floor", "2026-10-01", "--ceiling", "2026-11-12"], check=True, capture_output=True)


def check():
    """Every graded figure in the two files and in submission.md agrees with golden_figures.json."""
    f = json.loads((HERE / "golden_figures.json").read_text())
    bad = []


    def eq(a, b, what):
        if a != b:
            bad.append("%s: %r != %r" % (what, a, b))


    wb = load_workbook(TASK / "golden" / "trust_round_spring2027.xlsx", data_only=True)
    ws = wb["Spring round"]
    rows = [r for r in ws.iter_rows(min_row=5, values_only=True)]
    ten = [r for r in rows[:10]]
    eq([title_case(n) for n in f["round"]["slate"]], [r[1] for r in ten], "workbook slate")
    eq(f["round"]["planned"], [r[9] for r in ten], "workbook planned values")
    eq(f["round"]["running"], [r[10] for r in ten], "workbook running totals")
    eq([round(v) for v in f["round"]["new_recurring_revenue"]], [r[8] for r in ten], "workbook new recurring revenue")
    eq(f["round"]["total"], rows[10][9], "workbook round total")
    eq(f["round"]["running"][-1], f["round"]["total"], "the running total ends on the round")
    book = {(r[0], r[1]): r for r in wb["Book at 1 October"].iter_rows(min_row=5, values_only=True) if r[0]}
    for k, v in f["book"].items():
        t, b = k.split("|")
        r = book[(TYPE_LABEL[t], b)]
        want = (v["pupils"], v["revenue"]) if v["pupils"] else ("none", "none")
        eq(want, (r[2], r[3]), "workbook book " + k)
    pil = {r[0]: r for r in wb["Pilot licences"].iter_rows(min_row=5, values_only=True) if r[0]}
    for k, v in f["pilot"].items():
        r = pil[title_case(k)]
        eq((round(v["credit"]), round(v["net"])), (r[1], r[2]), "workbook pilot " + k)
    use = {(r[0], r[1]): r[2] for r in wb["Usage autumn 2026"].iter_rows(min_row=5, values_only=True) if r[0]}
    for k, v in f["usage"].items():
        t, b = k.split("|")
        eq(v if v is not None else "none", use[(TYPE_LABEL[t], b)], "workbook usage " + k)
    prs = Presentation(TASK / "golden" / "trust_round_spring2027.pptx")
    tbl = [sh for sh in prs.slides[0].shapes if sh.has_table][0].table
    eq([title_case(n) for n in f["round"]["slate"]], [tbl.cell(i, 1).text for i in range(1, 11)], "deck slate")
    title = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame][0]
    eq("£{:,}".format(f["round"]["total"]) in title, True, "deck title total")
    sub = (TASK / "submission.md").read_text()
    for i, (n, v) in enumerate(zip(f["round"]["slate"], f["round"]["running"]), 1):
        eq("%d. %s: £{:,}".format(v) % (i, title_case(n)) in sub, True, "submission running total " + n)
    b0 = f["round"]["behind"][0]
    eq("%s, £{:,} planned, £{:,} below the tenth".format(b0["planned"], f["round"]["gap"]) % title_case(b0["trust"])
       in sub, True, "submission first below the line")
    for k, v in f["pilot"].items():
        line = "%s: £{:,} credited, £{:,} net".format(round(v["credit"]), round(v["net"])) % title_case(k)
        eq(line in sub, True, "submission pilot " + k)
    for k, v in f["book"].items():
        t, b = k.split("|")
        if v["pupils"]:
            line = "%s, %s band: {:,} pupils, £{:,}".format(v["pupils"], v["revenue"]) % (TYPE_LABEL[t], b.lower())
            eq(line in sub, True, "submission book " + k)
    for k, v in f["usage"].items():
        t, b = k.split("|")
        if v is not None:
            eq("%s, %s band: %.1f" % (TYPE_LABEL[t], b.lower(), v) in sub, True, "submission usage " + k)
    r, e, cl = f["round"], f["round"]["eligible"], f["round"]["cluster"]
    for x in ["£{:,}".format(r["total"]), "£{:,}".format(r["gap"]), "**%d of %d**" % tuple(r["sign_rate"]),
              "all **12** pilot offers", "the sign rate of the %d Trust Licence offers closed to date" % r["sign_rate"][1],
              "**%d** of the %d open multi-academy trusts outside the pilot with two or more academies" % (e["central"], e["trusts"]),
              "%d open multi-academy trusts outside the pilot have at least two academies on agreements" % e["trusts"],
              "eligibility paragraph of trust_partnerships_plan_2027.docx: %d trusts" % e["central"],
              "Planned each of the %d" % e["central"],
              "Not %s, %s or the other trusts" % (title_case(r["rivals"][0]["trust"]), title_case(r["rivals"][1]["trust"]))]:
        eq(x in sub, True, "submission carries " + x)
    eq("—" in sub, False, "no em dash in submission")
    if bad:
        print("GOLDEN CHECK FAILED")
        for b in bad:
            print(" ", b)
        raise SystemExit(1)
    print("golden check: workbook, deck and submission agree with the figures")


if __name__ == "__main__":
    build()
    scrub()
    check()
    print("golden written to", OUT)
