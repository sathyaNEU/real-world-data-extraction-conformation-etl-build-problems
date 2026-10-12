"""The pack's documents: short, functional, in the voice of the people who wrote them."""
import datetime as dt

from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

COMPANY = "Tallybridge Maths Ltd"


def _pdf(path, title, blocks, author):
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=12.6,
                          alignment=TA_LEFT, spaceAfter=4)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=13, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10, spaceBefore=6,
                        spaceAfter=2)
    small = ParagraphStyle("s", parent=body, fontSize=7.8, textColor=colors.HexColor("#505050"))
    story = []
    for kind, val in blocks:
        if kind == "h1":
            story.append(Paragraph(val, h1))
        elif kind == "h2":
            story.append(Paragraph(val, h2))
        elif kind == "p":
            story.append(Paragraph(val, body))
        elif kind == "small":
            story.append(Paragraph(val, small))
        elif kind == "sp":
            story.append(Spacer(1, val * mm))
        elif kind == "table":
            t = Table(val, hAlign="LEFT")
            t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 8.8),
                                   ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.8),
                                   ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#333333")),
                                   ("LINEBELOW", (0, -1), (-1, -1), 0.4, colors.HexColor("#999999")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
            story.append(t)
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=22 * mm, rightMargin=22 * mm, topMargin=18 * mm,
                            bottomMargin=18 * mm, title=title, author=author, creator="", producer="")
    doc.build(story)


def licence_terms(path):
    _pdf(path, "Trust Licence terms", [
        ("h1", "Trust Licence: terms"),
        ("small", "Tallybridge Maths Ltd &middot; version 1.1 &middot; issued 12 May 2026 &middot; replaces version 1.0 "
                  "of 9 February 2026"),
        ("sp", 3),
        ("h2", "1. The parties"),
        ("p", "A Trust Licence is an agreement between Tallybridge Maths Ltd and an academy trust, for the academies set "
              "out in clause 2."),
        ("h2", "2. What it covers"),
        ("p", "Every academy the trust holds on the date the licence is priced, as Get Information about Schools (GIAS) "
              "shows it open in the trust on that date. An academy that joins the trust during a licence year is "
              "added at the next renewal. An academy that leaves the trust leaves the licence on the day it leaves."),
        ("h2", "3. Fee"),
        ("p", "&pound;4.05 for each pupil on roll for the licence year, excluding VAT, with no minimum per academy. "
              "Pupils on roll are the numbers GIAS publishes for each covered academy on the date the licence is "
              "priced."),
        ("h2", "4. Agreements the licence replaces"),
        ("p", "From the licence start date the licence replaces the Tallybridge agreement of every academy it covers. "
              "Where an academy has paid for part of its current term that falls after the licence start date, that "
              "part is credited to the academy."),
        ("h2", "5. Term"),
        ("p", "Twelve months from 1 September. The licence renews each year at the fee then current, on pupils as then "
              "published, unless either side gives notice before 31 May."),
        ("h2", "6. Invoicing"),
        ("p", "Annually in advance, to the trust. A trust that buys through a public buying framework is invoiced by "
              "the reseller appointed under that framework, and pays the reseller."),
        ("h2", "7. Earlier arrangements"),
        ("p", "The Trust Licence replaces the Cluster Licence, which was withdrawn on 31 August 2024. No Cluster "
              "Licence is in force."),
        ("sp", 4),
        ("small", "Commercial queries: trust partnerships team. Contract queries: finance. Registered in England and "
                  "Wales, company number 09884217."),
    ], "Tallybridge Maths Ltd")


def price_list(path):
    _pdf(path, "School subscriptions price list 2026-27", [
        ("h1", "School subscriptions: price list 2026-27"),
        ("small", "For terms starting on or after 1 April 2026. Prices exclude VAT, which is added at the standard rate."),
        ("sp", 3),
        ("table", [["", "Per pupil on roll, per year", "Minimum per agreement"],
                   ["Primary band", "£4.60", "£850"],
                   ["Secondary band", "£5.30", "£850"]]),
        ("sp", 3),
        ("p", "<b>Bands.</b> The primary band is for schools GIAS records as primary or middle deemed primary. Every "
              "other school is in the secondary band, including middle deemed secondary, all-through and special "
              "schools."),
        ("p", "<b>Pupils on roll.</b> The number GIAS publishes for the school when the term starts."),
        ("p", "<b>Split sites.</b> A school on more than one site may hold an agreement for each site, billed on the "
              "pupils taught there."),
        ("p", "<b>Terms that started before 1 April 2026</b> run to the end of the term at the 2025-26 prices: "
              "&pound;4.40 primary band, &pound;5.10 secondary band, &pound;820 minimum."),
        ("p", "<b>Trust Licence.</b> Separate terms apply; see the Trust Licence terms."),
    ], "Tallybridge Maths Ltd")


def plan_doc(path):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    d.add_heading("Trust partnerships: spring 2027 round", level=1)
    p = d.add_paragraph("Pamela Andrews, 8 October 2026. For the November planning meeting.")
    p.runs[0].italic = True
    d.add_heading("The round", level=2)
    d.add_paragraph("Leah O'Neill and Joel Bailey take the Trust Licence to five trusts each between the first week "
                    "of January and the Easter holidays, for licences from 1 September 2027. That is ten trusts. "
                    "Neither has room for an eleventh before Easter.")
    d.add_heading("Which trusts can be offered it", level=2)
    d.add_paragraph("An open multi-academy trust, as GIAS records it, in which at least two academies are on "
                    "Tallybridge agreements in force and the same person signed every one of those agreements. Trusts "
                    "offered the licence in the 2026 pilot, whether they signed or declined, are not offered it again "
                    "before 2028.")
    d.add_heading("How the ten are chosen", level=2)
    d.add_paragraph("The round goes to the trusts where it can win the most new recurring revenue.")
    d.add_heading("How a round goes on the plan", level=2)
    d.add_paragraph("Each trust's planned value is its new recurring revenue multiplied by the sign rate of the Trust "
                    "Licence offers closed to date, rounded to the whole pound. The round's value on the 2027-28 "
                    "plan is the sum of the ten planned values. Pupil numbers are taken as GIAS published them in the "
                    "extract of 7 October 2026, which this plan uses throughout; an academy with no published pupil "
                    "number counts nil.")
    d.add_heading("Timetable", level=2)
    for t in ["12 November: list signed off at the planning meeting.",
              "w/c 4 January: first approaches. Trust finance directors are hard to reach in the Easter holidays.",
              "Trusts set 2027-28 budgets over the spring and summer terms; most have them settled by July.",
              "By 31 July: licences signed for 1 September."]:
        d.add_paragraph(t, style="List Bullet")
    d.add_heading("Not in scope", level=2)
    d.add_paragraph("Local authority maintained schools and single-academy trusts stay on school agreements and are "
                    "handled by the account managers as now.")
    d.save(path)


def data_notes(path):
    text = """# Tallybridge exports: field notes

Maintained by revenue operations (Maureen Webb). Last updated 19 October 2026.

These notes cover the exports revenue operations produces. GIAS files are as DfE publishes them and are described by DfE.

## tb_agreements_2026-10-01.csv

One row per agreement, current and past, as held on 1 October 2026. An agreement is identified by `agreement_ref`.

| field | meaning |
|---|---|
| agreement_ref | the agreement's reference |
| account_id, account_name | the CRM account the agreement sits on, and its name as recorded |
| urn | the school's URN as recorded when its account was created. Blank for trust accounts |
| dfe_number | LA code and establishment number, as recorded when the account was created |
| companies_house_no | the trust's company number, trust accounts only |
| billing_basis | Per pupil, Trust Licence or Cluster Licence |
| price_band | the price list band the agreement is billed in |
| pupils_billed | the pupils on roll the CRM holds against the current (or last) term |
| annual_value_gbp | the annual value the CRM holds against the current (or last) term, excluding VAT |
| agreement_start | when the agreement first started |
| term_start, term_end | the current term, or the last term for an agreement that has ended |
| end_date | the last day of an agreement that has ended or been replaced |
| status | In force, Ended, or Replaced by Trust Licence |
| signed_by, signed_by_role, signed_on | who signed the agreement in its current form, and when |
| signed_by_contact_id | the CRM contact who signed |

## tb_crm_opportunities_2022-2026.csv

One row per opportunity in the CRM from January 2022. `opportunity_type` is New business, Renewal, Cluster Licence or Trust Licence. `amount_gbp` is the annual value quoted. `lost_reason` is the CRM picklist value chosen when an opportunity is closed lost. Opportunities for schools that never became customers carry no account.

## tb_billing_ledger_2025-04_2026-09.csv

One row per invoice or credit note issued between 1 April 2025 and 30 September 2026, from both billing systems. `amount_gbp` is the amount on the document as the issuing system holds it and `vat_gbp` is the VAT that system holds against it. For a credit note, `reference` is the document it credits. Documents issued through the reseller are not in this ledger.

## tb_weekly_usage_2025-09_2026-10.parquet

One row per school per reporting week in which at least one of its pupils answered a question. `school_ref` is the platform's school id and `urn` is the URN on the platform record. `active_pupils` counts distinct pupils who answered at least one question in the week. Reporting weeks run Monday to Sunday and each weekly load reports one of them: `batch_id` names the load by year and ISO week number, and `loaded_at` is when it was loaded. `week_start` is the date the reporting service labels the week with. A week sent again for some schools is loaded under that week's `batch_id` with a suffix, and the later load is the one of record for the schools it carries.

## trust_licence_pilot_2026.xlsx

The twelve Trust Licence offers made in the 2026 pilot, priced on 15 June 2026, with the schedule of academies for each licence signed.

## orrin_remittance_2026Q3.csv

The reseller's statement for July to September 2026: a line for each school it supplies under the public buying framework, with `contract_ref`, the reseller's reference for the contract the school is supplied under, and a total line.

## Definitions used in our reporting

- **The book**: what Tallybridge holds on a date: agreements in force, Trust Licences in force and contracts held through the reseller.
- **Pupils billed**: the pupils on roll a term is invoiced on.
- **Recurring revenue**: what an agreement or licence brings Tallybridge in a year, excluding VAT, as the term in force is invoiced. For a contract held through the reseller it is the annual contract value less the reseller's margin.
- **Weekly active pupils per hundred billed**: the mean of a group's weekly active pupils over the weeks taken, per hundred pupils billed.
- **Autumn term weeks**: counted from the week beginning Monday 7 September 2026.
- **Rounding**: figures are carried in pence and rounded once, at the figure reported.
"""
    open(path, "w").write(text)


def thread(path):
    text = """#trust-partnerships  (export, 20 Oct 2026)

**Henry Miles** 14 Oct, 09:12
Pamela, can you bring the spring round to the planning meeting on the 12th? Ten trusts, in the order Leah and Joel should go to them, and what the round puts on next year's plan.

**Pamela Andrews** 14 Oct, 09:31
Will do. Plan note is in the folder, I'll work from the 7 Oct GIAS extract.

**Philip Saunders** 14 Oct, 10:02
My two pence: go where the room is. The big trusts where we've only got two or three schools are where the money is, not the ones we've already sold to.

**Maureen Webb** 14 Oct, 10:15
The pilot sheet has every quote and booking from the summer if you want to check your numbers against it.

**Joel Bailey** 14 Oct, 11:48
Easter holidays are 29 Mar to 9 Apr for most of them. Finance directors vanish, so we're really looking at Jan to mid-March.

**Henry Miles** 16 Oct, 08:47
Also want the base on the same deck: what we hold by type of school, how the pilot licences have billed so far, and whether schools are actually using it this term.

**Pamela Andrews** 16 Oct, 09:05
OK. Usage export is refreshed weekly, I'll take the first six weeks of term.

**Leah O'Neill** 16 Oct, 09:20
:+1:
"""
    open(path, "w").write(text)


def finance_note(path):
    text = """From: Maureen Webb
To: Finance; Revenue Operations; Account Management
Date: Fri, 24 Apr 2026 16:12
Subject: Billing platform from 1 May

Hi all,

From Friday 1 May every new invoice and credit note comes from the new billing platform. Its documents are numbered INV- and CN-.

The old system stays open until 31 March 2027, but only for credit notes against its own invoices. So if we credit anything that was invoiced before 1 May, the credit note still comes out of the old system with an L number, same as the invoice did.

Nothing changes for framework schools: the reseller invoices those and we see them on its quarterly remittance.

Thanks,
Maureen
"""
    open(path, "w").write(text)


def sources(path, counts):
    text = """About this folder: where each file comes from and the terms it is held under
Prepared by revenue operations, 20 October 2026.

Department for Education, Get Information about Schools (GIAS), public downloads of 7 October 2026, as published:
  edubasealldata20261007.csv        establishments, all statuses ({est} rows)
  links_edubasealldata20261007.csv  links between establishments ({links} rows)
  allgroupsdata20261007.csv         groups: trusts, federations, sponsors ({groups} rows)
  Source: https://get-information-schools.service.gov.uk/Downloads
  Licence: Open Government Licence v3.0. Contains public sector information licensed under the Open Government Licence v3.0.

Tallybridge Maths Ltd internal exports (not for circulation outside Tallybridge):
  tb_agreements_2026-10-01.csv              CRM agreements, as at 1 October 2026 ({book} rows)
  tb_crm_opportunities_2022-2026.csv        CRM opportunities from January 2022, pulled 19 October 2026 ({crm} rows)
  tb_billing_ledger_2025-04_2026-09.csv     billing ledger, both billing systems, pulled 5 October 2026 ({ledger} rows)
  tb_weekly_usage_2025-09_2026-10.parquet   platform weekly usage to the week beginning 12 October 2026 ({usage} rows)
  trust_licence_pilot_2026.xlsx             trust partnerships team, pilot offers and licence schedules
  tallybridge_trust_licence_terms_2026.pdf  version 1.1
  tallybridge_price_list_2026-27.pdf
  trust_partnerships_plan_2027.docx
  revenue_ops_export_guide.md               field notes on the exports above
  trust_round_thread.md                     channel export
  finance_note_billing_cutover.txt          email, 24 April 2026

Third party:
  orrin_remittance_2026Q3.csv  Orrin Educational Supply Ltd, quarterly remittance under the public buying framework,
                               July to September 2026 ({remit} lines and a total line). Supplied to Tallybridge for
                               reconciliation; not to be passed on.
"""
    open(path, "w").write(text.format(**counts))
