"""task153 checks: the three rules executed on the clean truth, the shipped HOW sentences executed literally, the
ladder and every rival route, the corridor of defensible golden constructions, every semantic fork row, the closure
axes that can move, the 2025 out-of-sample record, both asks with their stops and devices, the separation line, the
realism sweep, the input gates and the distractors. Run by build.py; exits on the first report with any failure."""
from __future__ import annotations

import collections
import copy
import datetime as dt
import json
import math
import os
import re
import shutil
import statistics
import tempfile

import openpyxl
import pdfplumber

import asks as AS
import calc as CA
import common as K
import deliver as DL
import docs as DOC
import records as R
import world as WD

RESULTS = []
# the design's figures, asserted as built (a parameter change that moves one fails here first)
EXPECT = {
    "R2b": 12, "SF3": 28,
    "A": {("West", (2026, 5)): 33472, ("West", (2026, 6)): 36732, ("West", (2026, 7)): 37170,
          ("Central", (2026, 5)): 46569, ("Central", (2026, 6)): 54781, ("Central", (2026, 7)): 52821,
          ("East", (2026, 5)): 46391, ("East", (2026, 6)): 51215, ("East", (2026, 7)): 53258},
    "B": {("JXN", (2025, 3)): (31911, 891), ("JXN", (2025, 4)): (24004, 811), ("JXN", (2025, 5)): (7986, 254),
          ("KNX", (2025, 6)): (23977, 446), ("KNX", (2025, 7)): (24021, 542), ("KNX", (2025, 8)): (8021, 191)},
}
GOLD = {"Memphis": 2, "Jackson": 5, "Nashville": 8, "Cookeville": 1, "Knoxville": 4, "Blountville": 3}


def ok(label, cond, detail=""):
    RESULTS.append((label, bool(cond), detail))


def counts(d):
    return {y: CA.floor_loads(d[y]) for y in CA.YARDS}


def tot(d):
    return sum(counts(d).values())


def name(code):
    return K.YARD[code]["name"]


def world_routes(W, r_hat):
    """The easy route and the rungs on the generator's clean truth: the world's card fills by pump date, its legs'
    fuel stretch by stretch between yard arrivals (or between nights at the yard), each unit's true tankful above
    its driver's reserve, and the share the 2025 islands handed out."""
    a, b = CA.BASE
    card, visit, over, fuel = (collections.defaultdict(float) for _ in range(4))
    for u in W.units:
        if u.status != "Active":
            continue
        y = name(u.yard)
        for f in u.fills:
            if f["kind"] == "card" and a <= f["t"].date() <= b:
                card[y] += f["gal"]
        cap = u.cap - u.reserve
        acc = acc_o = 0.0
        for i, leg in enumerate(u.legs):
            acc += leg[4]
            acc_o += leg[4]
            if a <= leg[0].date() <= b:
                fuel[y] += leg[4]
            if leg[3] == ("YARD", u.yard):
                t1 = leg[2]
                inb = a <= t1.date() <= b
                if inb:
                    visit[y] += min(acc, cap)
                acc = 0.0
                nxt = u.legs[i + 1][0] if i + 1 < len(u.legs) else None
                if nxt is None or nxt.date() > t1.date():
                    if inb:
                        over[y] += min(acc_o, cap)
                    acc_o = 0.0
    k = CA.working_days(*CA.WIN) / CA.working_days(*CA.BASE)
    isl = collections.defaultdict(float)
    burn = collections.defaultdict(float)
    for u in W.units:
        if u.yard not in K.PROG:
            continue
        lo, hi = WD.ISLAND[u.yard]
        isl[u.yard] += sum(f["gal"] for f in u.fills if f["kind"] == "island")
        burn[u.yard] += sum(l[4] for l in u.legs if lo <= l[0].date() <= hi)
    share = sum(isl.values()) / sum(burn.values())
    return dict(card={y: card[y] for y in CA.YARDS}, visit={y: visit[y] * k for y in CA.YARDS},
                over={y: over[y] * k for y in CA.YARDS}, fuel={y: fuel[y] for y in CA.YARDS},
                share=share, share_yard={y: isl[y] / burn[y] for y in isl})


def source_text():
    """The documents' text as written: the PDFs from their source strings, the markdown files in full."""
    parts = list(DOC.OFFER_OPEN) + ["%s: %s" % kv for kv in DOC.OFFER_TERMS] + list(DOC.OFFER_CLOSE)
    for L in DOC.PROG_LETTERS.values():
        parts += L["body"] + ["%s: %s" % kv for kv in L["terms"]] + L["close"]
    for c in DOC.CARD_SCHEDULES:
        parts += [c["title"], c["party"], c["signed"]] + c["clauses"]
    for k, rows in DOC.LAYOUT.items():
        parts += [" ".join(r) for r in rows]
    parts += DOC.LAYOUT_NOTES + [" ".join(r) for r in DOC.TW_ROWS] + DOC.TW_NOTES
    return "\n".join(parts)


def pdf_text(path):
    with pdfplumber.open(path) as pdf:
        return "\n".join(pg.extract_text() or "" for pg in pdf.pages)


def run(target, golden, B=None, G=None, meta_path=None):
    W, X = B.W, G
    P, Gd = X["P"], X["G"]
    draw, loads = Gd["draw"], Gd["loads"]
    Rg, S = X["R"], X["S"]

    # ------------------------------------------------------------------ the call
    ok("CALL 23 loads: Memphis 2, Jackson 5, Nashville 8, Cookeville 1, Knoxville 4, Blountville 3",
       loads == GOLD and Gd["total"] == 23, str(loads))
    clear = {y: K.floor_clearance(draw[y], CA.LOAD) for y in CA.YARDS}
    ok("CALL every yard's draw sits inside its load bin by 0.25 of a load or more", min(clear.values()) >= 0.25,
       ", ".join("%s %.2f" % (CA.CODE[y], draw[y] / CA.LOAD) for y in CA.YARDS))
    ok("CALL saving $17,408: 23 loads of 8,000 gallons at each account's charge over the TriState average less $0.019",
       abs(X["saving"] - 17408) < 0.01 and X["per_load"] == {"West": 800.0, "Central": 640.0, "East": 864.0},
       "%.2f" % X["saving"])
    big = [f for f in P.fms if f["gal"] > 225]
    ok("CALL the reserve read from the 2025 islands: every fill over 225 gallons is a T680 sleeper's, %d fills "
       "averaging %.1f of 300, and no T680 fill falls between 215 and 255 gallons" % (Gd["n_reserve"],
                                                                                      300 * (1 - Gd["reserve"])),
       len(big) == Gd["n_reserve"] and all(P.units[f["unit"]]["model"] == "T68SL" for f in big) and
       not any(215 < f["gal"] < 255 for f in P.fms if P.units[f["unit"]]["model"] == "T68SL") and
       265 < 300 * (1 - Gd["reserve"]) < 273, "%.4f" % Gd["reserve"])

    # the corridor: every defensible construction of the golden lands the same six counts
    variants = {}
    variants["raw eight weeks (no working-day scaling)"] = CA.draw_by_yard(P, scale=False)
    full8 = [(dt.date(2026, 6, 22), dt.date(2026, 6, 28)), (dt.date(2026, 7, 6), CA.BASE[1])]
    v = collections.defaultdict(float)
    for (a, b) in full8:
        for y, g in CA.draw_by_yard(P, base=(a, b), scale=False).items():
            v[y] += g
    variants["the eight full working weeks (the holiday week left out)"] = dict(v)
    b13 = (CA.BASE[1] - dt.timedelta(days=90), CA.BASE[1])
    variants["13 weeks, scaled by working days"] = CA.draw_by_yard(P, base=b13)
    variants["13 weeks, scaled by weeks"] = {y: g * 8 / 13 for y, g in CA.draw_by_yard(P, base=b13, scale=False).items()}
    b25 = (dt.date(2026, 3, 2), CA.BASE[1])
    variants["25 weeks from 2 March, scaled by working days"] = CA.draw_by_yard(P, base=b25)
    variants["the nominal tank (no reserve)"] = CA.draw_by_yard(P, cap_mode="nominal")
    maxfill = collections.defaultdict(float)
    for f in P.fms:
        maxfill[f["unit"]] = max(maxfill[f["unit"]], f["gal"])
    caps = {u: maxfill[u] for u in maxfill if P.units[u]["model"] == "T68SL" and maxfill[u] > 225}
    variants["each 2025 sleeper's own largest island fill as its tankful"] = CA.draw_by_yard(P, caps=caps)
    variants["a midday visit's day split in equal halves"] = CA.draw_by_yard(P, split="equal")
    # the season: the 2025 fuel per working day in the window's weeks against the base's weeks
    w25 = (dt.date(2025, 9, 15), dt.date(2025, 11, 9))
    b25w = (dt.date(2025, 6, 30), dt.date(2025, 8, 24))
    f_w = sum(r["fuel"] for u in P.days for d, r in P.days[u].items() if w25[0] <= d <= w25[1])
    f_b = sum(r["fuel"] for u in P.days for d, r in P.days[u].items() if b25w[0] <= d <= b25w[1])
    season = (f_w / CA.working_days(*w25)) / (f_b / CA.working_days(*b25w))
    variants["the 2025 seasonal ratio applied (%.3f)" % season] = {y: g * season for y, g in draw.items()}
    variants["13 weeks by weeks, with the seasonal ratio"] = {
        y: g * season for y, g in variants["13 weeks, scaled by weeks"].items()}
    variants["raw eight weeks, with the seasonal ratio"] = {
        y: g * season for y, g in variants["raw eight weeks (no working-day scaling)"].items()}
    # engine fuel calibrated to pump gallons unit by unit over two years (cards and islands)
    pump = collections.defaultdict(float)
    for f in P.cards:
        if f["diesel"] and f["unit"] in P.units:
            pump[f["unit"]] += f["gal"]
    for f in P.fms:
        pump[f["unit"]] += f["gal"]
    ecu = {u: sum(r["fuel"] for r in P.days[u].values()) for u in P.days}
    fscale = {u: pump[u] / ecu[u] for u in ecu if ecu[u] > 0 and pump[u] > 0}
    variants["engine fuel calibrated to pump gallons by unit"] = CA.draw_by_yard(P, fuel_scale=fscale)
    bad = {k: counts(v) for k, v in variants.items() if counts(v) != GOLD}
    ok("CALL every defensible construction lands the same six counts (%d constructions)" % len(variants), not bad,
       str(bad))
    lo = {y: min(v[y] for v in variants.values()) / CA.LOAD for y in CA.YARDS}
    hi = {y: max(v[y] for v in variants.values()) / CA.LOAD for y in CA.YARDS}
    ok("CALL across the constructions each yard stays 0.15 of a load or more inside its bin",
       all(lo[y] - GOLD[y] >= 0.15 and GOLD[y] + 1 - hi[y] >= 0.15 for y in CA.YARDS),
       ", ".join("%s %.2f-%.2f" % (CA.CODE[y], lo[y], hi[y]) for y in CA.YARDS))
    # one sigma: the eight-week total's sampling spread from the base's regular weeks
    weekly = collections.defaultdict(lambda: collections.defaultdict(float))
    r_hat = Gd["reserve"]
    for u, info in P.units.items():
        if info["status"] != "Active":
            continue
        cap = info["tank"] * (1 - r_hat)
        for (a, g) in CA.island_fills(P, u, cap):
            if CA.BASE[0] <= a.date() <= CA.BASE[1]:
                weekly[info["yard"]][a.date() - dt.timedelta(days=a.weekday())] += g
    sig = {}
    for y in CA.YARDS:
        vals = [g for wk, g in weekly[y].items() if wk != dt.date(2026, 6, 29)]
        sig[y] = statistics.stdev(vals) * math.sqrt(8) * 40 / 39
    ok("CALL each yard's count holds one sampling sigma either way",
       all(CA.floor_loads(draw[y] - sig[y]) == GOLD[y] == CA.floor_loads(draw[y] + sig[y]) for y in CA.YARDS),
       ", ".join("%s +-%.0f" % (CA.CODE[y], sig[y]) for y in CA.YARDS))

    # ------------------------------------------------------------------ the three rules on the clean truth
    nb = sum(1 for u in P.days for d in P.days[u] if CA.BASE[0] <= d <= CA.BASE[1])
    nv = X["build"]
    ok("R1a the decisive computation reads the vehicle days and pairs the geofence crossings into yard visits",
       nb >= 5000 and sum(b["visits"] for b in nv.values()) >= 4000,
       "%d vehicle days in the base, %d yard visits" % (nb, sum(b["visits"] for b in nv.values())))
    heads = []
    for f in sorted(os.listdir(target)):
        p = os.path.join(target, f)
        if f.endswith(".csv"):
            with open(p, encoding="utf-8") as fh:
                heads.append((f, fh.readline()))
        elif f.endswith(".xlsx"):
            wb = openpyxl.load_workbook(p, read_only=True)
            for ws in wb.worksheets:
                heads.append((f, " ".join(str(c) for c in next(ws.iter_rows(values_only=True)))))
    hit = [f for f, h in heads if re.search(r"draw|captur|stretch|visit|convertib|share|segment|reserve", h, re.I)]
    ok("R1b no shipped column holds an island draw, a capture share, a stretch, a visit count or a reserve", not hit,
       str(hit))
    # R2: each Layer 1 item mishandled leaves the six counts where they are
    P_nodev = CA.read(target, cards=False, use_devices=False)
    ok("R2a the new devices' rows under their serials dropped (the assignment table not used): same counts",
       counts(CA.draw_by_yard(P_nodev)) == GOLD)
    P_dup = CA.read(target, cards=False, dedupe_fms=False)
    ok("R2b the island batch resent after the Knoxville modem outage counted twice: same counts",
       P_dup.fms_dupes == 14 and counts(CA.draw_by_yard(P_dup)) == GOLD, "%d resent rows" % P_dup.fms_dupes)
    P_utc = copy.copy(P)
    P_utc.stays = {}
    for u, Sx in P.stays.items():
        tz = CA.TZ[P.units[u]["yard"]] if u in P.units else CA.CT
        P_utc.stays[u] = [((a.replace(tzinfo=tz) if a else None), (b.replace(tzinfo=tz) if b else None), g)
                          for (a, b, g) in Sx]
    ok("R2c the geofence export's UTC stamps read as local time: same counts", counts(CA.draw_by_yard(P_utc)) == GOLD,
       str(counts(CA.draw_by_yard(P_utc))))
    P_ct = CA.read(target, cards=False, day_tz=CA.CT)
    ok("R2d the vehicle days' Eastern clock read as Central: same counts", counts(CA.draw_by_yard(P_ct)) == GOLD)
    P_cls = copy.copy(P)
    P_cls.units = {u: dict(i, tank=300.0 if i["cls"] == "Tractor" else i["tank"]) for u, i in P.units.items()}
    ok("R2e every tractor given the sleeper's tank (the spec joined by class, not model): same counts",
       counts(CA.draw_by_yard(P_cls)) == GOLD)
    P_nocard = CA.read(target, cards=False)
    ok("R2f the card exports (unit numbers written three ways, three layouts) are not on the golden's path",
       counts(CA.draw_by_yard(P_nocard)) == GOLD and not hasattr(P_nocard, "cards"))
    T = world_routes(W, r_hat)
    ok("R3a the analyst test on the clean truth: the islands taking over the vehicles' card gallons file 29",
       tot(T["card"]) == 29 and counts(T["card"]) == {"Memphis": 2, "Jackson": 5, "Nashville": 8, "Cookeville": 3,
                                                      "Knoxville": 5, "Blountville": 6}, str(counts(T["card"])))
    ok("R3b the decisive step alone on the clean truth (each stretch between yard visits capped at the true tankful) "
       "files 23 with the golden's six counts", counts(T["visit"]) == GOLD, str(counts(T["visit"])))
    ok("R3b each lower rung's single correction on the clean truth stays wrong: the 2025 share 26, the overnight cut 19",
       sum(CA.floor_loads(T["card"][y] * T["share"]) for y in CA.YARDS) == 26 and tot(T["over"]) == 19,
       "share %.3f, overnight %s" % (T["share"], counts(T["over"])))
    rep = 0
    seen = set()
    for u in P.days:
        ds = sorted(P.days[u])
        for i in range(len(ds) - 4):
            key = tuple(P.days[u][d]["fuel"] for d in ds[i:i + 5])
            if key in seen:
                rep += 1
            seen.add(key)
    ok("R3c no copied block: no five-day run of a vehicle's fuel repeats anywhere in the vehicle days", rep == 0, rep)
    texts = {}
    for f in DOC.DOCS:
        p = os.path.join(target, f)
        texts[f] = pdf_text(p) if f.endswith(".pdf") else open(p, encoding="utf-8").read()
    alltext = "\n".join(texts.values())
    src = source_text() + "\n" + texts[DOC.F_NOTES] + "\n" + texts[DOC.F_README]
    plants = re.findall(r"\b\d+\s+loads?\b|\bloads? of \d|\bdraw\b|\bcapture|\bconvert|\bshare of\b", src, re.I)
    ok("R3c no document states a load count for the window, an island draw or a capture share", not plants, plants)

    # ------------------------------------------------------------------ the shipped HOW sentences, executed literally
    how = {
        "offer: load": "8,000 net gallons (corrected to 60 degrees F), lifted for and delivered to one yard tank.",
        "offer: count": "Up to 32 in all, for any of the six yards.",
        "offer: commitment": "Every confirmed load is billed at the price above whether or not it has been called off "
                             "by November 8.",
        "offer: price": "The TriState Rack Report contract average for ULSD No. 2 at the supplying terminal on the day "
                        "of lifting, the same average your three card programmes bill from, plus $0.019 per net "
                        "gallon, delivered.",
        "notes: days": "Every yard runs Monday to Friday",
        "notes: fuel": "Each account's vehicles carry that account's cards and buy wherever the route takes them.",
    }
    flat = re.sub(r"\s+", " ", src)
    missing = [k for k, v in how.items() if re.sub(r"\s+", " ", v) not in flat]
    ok("HOW the six shipped sentences on how loads, gallons and prices are counted are quoted from the pack",
       not missing, missing)
    ptxt = open(os.path.join(os.path.dirname(os.path.abspath(target)), "prompt.txt"), encoding="utf-8").read()
    ok("HOW the prompt files the scenario and both chart windows: our own vehicles running as in the eight weeks to 23 "
       "August, the draw for the contract weeks, the cards' diesel of those eight weeks",
       ptxt.count("the eight weeks to 23 August") == 2 and "if they run as in the eight weeks to 23 August" in ptxt and
       "island draw for the contract weeks" in ptxt and CA.BASE == (dt.date(2026, 6, 29), dt.date(2026, 8, 23)))
    lit = {y: min(Rg["R0 card gallons, trailing eight weeks"][y], 32 * CA.LOAD) for y in CA.YARDS}
    ok("HOW executed word for word on the gallons the files record as bought (whole loads per yard tank, 32 at most, "
       "Monday-to-Friday weeks replayed with no change) they file 29, never the golden",
       tot(lit) == 29 and counts(lit) != GOLD, str(counts(lit)))
    cue = re.findall(r"tankful|reserve|midday|mid-day|wednesday|yard visit|between visits|reload|top off|top-off|"
                     r"once a week|every monday|inspection", src, re.I)
    ok("HOW no shipped sentence says when an island can fuel a vehicle or how much one fill puts aboard", not cue, cue)

    # ------------------------------------------------------------------ the ladder
    R0 = Rg["R0 card gallons, trailing eight weeks"]
    ok("LADDER R0 the card gallons of the trailing eight weeks fill 29 (2, 5, 8, 3, 5, 6)",
       tot(R0) == 29 and counts(R0) == {"Memphis": 2, "Jackson": 5, "Nashville": 8, "Cookeville": 3, "Knoxville": 5,
                                        "Blountville": 6})
    fam = ["R0 card gallons x 40/39", "R0 card gallons, 13 weeks", "R0 card gallons, same weeks of 2025"]
    ok("LADDER R0's family (scaled, 13 weeks, the same weeks of 2025) files 29 on every base",
       all(tot(Rg[k]) == 29 for k in fam), str([tot(Rg[k]) for k in fam]))
    r1 = ["R1a pooled 2025 share", "R1a' 2025 share by account", "R1b 2025 share by model"]
    r1s = [k + " x 40/39" for k in r1]
    ok("LADDER R1 the 2025 islands' share carried onto the card gallons, pooled, by account or by model, files 26 on "
       "the eight weeks as bought and 26 to 28 scaled to the window's working days (pooled 28, by account 27, by "
       "model 26)", all(tot(Rg[k]) == 26 for k in r1) and [tot(Rg[k]) for k in r1s] == [28, 27, 26],
       str([tot(Rg[k]) for k in r1 + r1s]))
    R2a = Rg["R2a tankful per run between nights at the yard"]
    ok("LADDER R2a a tankful per run between nights at the yard files 19 (2, 5, 5, 0, 4, 3)",
       tot(R2a) == 19 and counts(R2a) == {"Memphis": 2, "Jackson": 5, "Nashville": 5, "Cookeville": 0,
                                          "Knoxville": 4, "Blountville": 3})
    ok("LADDER R2b the card gallons times each vehicle's share of nights at the yard files %d" % EXPECT["R2b"],
       tot(Rg["R2b card gallons x share of nights at the yard"]) == EXPECT["R2b"])
    refl = ["Reflex card fills in the yard's own city", "Reflex daily-return vehicles only",
            "Reflex mean of the card gallons and the overnight cut"]
    ok("LADDER the reflexes (fills in the yard's city 6, daily-return vehicles 7, the mean of R0 and R2a 24) each "
       "miss the total and two yards or more",
       [tot(Rg[k]) for k in refl] == [6, 7, 24] and all(
           sum(counts(Rg[k])[y] != GOLD[y] for y in CA.YARDS) >= 2 for k in refl), str([counts(Rg[k]) for k in refl]))
    every = [k for k in Rg if k != "Golden"]
    ok("LADDER no rung or reflex files 23, and each differs from the golden at one yard or more",
       all(tot(Rg[k]) != 23 and counts(Rg[k]) != GOLD for k in every))
    ok("LADDER the golden is bracketed: the card gallons and the shares above, the overnight cut below; only the cut at "
       "every yard visit reverses it (Nashville 5 to 8, Cookeville 0 to 1)",
       tot(R0) > tot(Rg["R1a pooled 2025 share"]) > 23 > tot(R2a) and
       counts(R2a)["Nashville"] == 5 and counts(R2a)["Cookeville"] == 0)
    # the 2025 record, out of sample
    spans = S["span"]
    pre = {}
    for code in ("JXN", "KNX"):
        a = spans[code][0]
        pa = a - dt.timedelta(days=56)
        cu = CA.card_by_unit(P, pa, a - dt.timedelta(days=1))
        units = [u for u, i in P.units.items() if CA.CODE[i["yard"]] == code]
        k_ = CA.working_days(*spans[code]) / CA.working_days(pa, a - dt.timedelta(days=1))
        pre[code] = sum(cu[u] for u in units) * k_
    isl = {c: sum(f["gal"] for f in P.fms if f["site"] == c) for c in ("JXN", "KNX")}
    err0 = {c: pre[c] / isl[c] - 1 for c in isl}
    ok("OOS the 2025 islands refute the card gallons: the vehicles' card gallons of the eight weeks before overstate "
       "what each island handed out (Knoxville by 10 per cent or more)", err0["KNX"] >= 0.10 and err0["JXN"] > 0,
       "JXN %+.3f, KNX %+.3f" % (err0["JXN"], err0["KNX"]))
    ok("OOS each 2025 island's share misses the other's: Jackson's carried to Knoxville, or Knoxville's to Jackson, "
       "off by 8 per cent or more", abs(S["yard"]["Jackson"] / S["yard"]["Knoxville"] - 1) >= 0.08,
       "JXN %.3f, KNX %.3f" % (S["yard"]["Jackson"], S["yard"]["Knoxville"]))
    bt = {}
    for code in ("JXN", "KNX"):
        a, b = spans[code]
        ysum = 0.0
        for u, i in P.units.items():
            if CA.CODE[i["yard"]] != code or i["status"] != "Active":
                continue
            cap = i["tank"] * (1 - r_hat)
            acc = 0.0
            for d in sorted(P.days.get(u, {})):
                if not (a <= d <= b):
                    continue
                rr = P.days[u][d]
                acc += rr["fuel"]
                if CA.near_yard(P, rr["lat1"], rr["lon1"]) == i["yard"]:
                    ysum += min(acc, cap)
                    acc = 0.0
        bt[code] = ysum / isl[code] - 1
    ok("OOS a tankful per run between nights at the yard, from the vehicle days alone, reproduces each 2025 island "
       "within 3 per cent (the record cannot tell that cut from the visit cut: no 2025 run that outlasted a tankful "
       "came back during the day)", all(abs(v) <= 0.03 for v in bt.values()),
       "JXN %+.3f, KNX %+.3f" % (bt["JXN"], bt["KNX"]))
    jmid = sum(1 for f in P.fms if f["site"] == "JXN" and P.units[f["unit"]]["model"] == "CAS1D" and
               11 <= f["t"].hour <= 13)
    ok("OOS the 2025 record shows an island fill at a midday yard visit (Jackson's day cabs on their midday reloads)",
       jmid >= 100, "%d midday day-cab fills" % jmid)

    # ------------------------------------------------------------------ the semantic fork table
    gross = {y: draw[y] / 1.0055 for y in CA.YARDS}
    ok("SF1 a load read as 8,000 gallons at the island meter (gross, 0.55 per cent over net): same call",
       counts(gross) == GOLD)
    room = {y: draw[y] + (K.TANKS[CA.CODE[y]][1] - K.TANKS[CA.CODE[y]][2]) for y in CA.YARDS}
    ok("SF2 loads delivered into the tank by 8 November (not handed out) would file 30 or more", min(32, tot(room)) >= 30)
    near = sum(int(math.floor(draw[y] / CA.LOAD + 0.5)) for y in CA.YARDS)
    ok("SF3 the nearest whole load instead of the whole loads within the draw files %d" % EXPECT["SF3"],
       near == EXPECT["SF3"], near)
    ok("SF4 the fleet's total draw floored once files 26 (each load is one drop at one yard tank)",
       CA.floor_loads(sum(draw.values())) == 26)
    other = 0
    for u, Sx in P.stays.items():
        own = "KSC %s Yard" % P.units[u]["yard"]
        other += sum(1 for (a, b, g) in Sx if g.endswith(" Yard") and g != own)
    ok("SF5 no unit enters another yard's geofence in 2026: 'our own vehicles' has one reading", other == 0, other)
    nockv = dict(draw)
    vans = sum(g for u, i in P.units.items() if i["yard"] == "Cookeville" and i["model"] == "SPR25"
               for (a, g) in CA.island_fills(P, u, i["tank"] * (1 - r_hat)) if CA.BASE[0] <= a.date() <= CA.BASE[1])
    nockv["Cookeville"] = draw["Cookeville"] - vans * 40 / 39
    ok("SF6 Cookeville's lot-based vans left out would file 22 (the unit list puts them at Cookeville)",
       tot(nockv) == 22 and all(P.units[u]["yard"] == "Cookeville" for u, i in P.units.items() if
                                i["model"] == "SPR25" and str(u).startswith("35")))
    wkend = sum(g for u, i in P.units.items() if i["status"] == "Active"
                for (a, g) in CA.island_fills(P, u, i["tank"] * (1 - r_hat))
                if CA.BASE[0] <= a.date() <= CA.BASE[1] and a.weekday() >= 5)
    ok("SF7 no vehicle reaches its yard on a weekend: the window read as 56 days or as 40 working days is the same",
       wkend == 0)
    ok("SF8 the base period: 8, 13 or 25 weeks, scaled by working days or by weeks, with or without the season: same "
       "call (the corridor above)", not bad)
    ok("SF9 the tankful: the tank less the 2025 reserve, the nominal tank, or each sleeper's largest 2025 fill: same "
       "call", counts(variants["the nominal tank (no reserve)"]) == GOLD and counts(
           variants["each 2025 sleeper's own largest island fill as its tankful"]) == GOLD)
    nofee = X["loads_acct"]["West"] * 8000 * 0.100 + X["loads_acct"]["Central"] * 8000 * (0.084 - 0.019) + \
        X["loads_acct"]["East"] * 8000 * 0.108
    ok("SF10 the Crossroads network fee is part of its diesel price (Addendum 4 clause 2): $0.080 a gallon saved, not "
       "the $0.065 without it ($%s in all, not the agreement terms' saving)" % format(round(nofee), ","),
       round(nofee) == 16328 and round(nofee) != round(X["saving"]) and X["per_gal"]["Central"] == 0.08 and
       "is part of the diesel price under this addendum" in flat)

    rack = {}
    import csv as _csv
    with open(os.path.join(target, R.F_RACK), encoding="utf-8") as fh:
        for r in _csv.DictReader(fh):
            rack[(r["Terminal"], dt.date.fromisoformat(r["Effective Date"]))] = float(r["TriState Contract Avg"])
    devs = []
    with open(os.path.join(target, R.F_WEST), encoding="utf-8") as fh:
        for r in _csv.DictReader(fh):
            if r["PROD_CD"] != "01" or r["REF_TRAN_ID"]:
                continue
            t = dt.datetime.strptime(r["TRAN_DATE_TIME"], "%Y-%m-%d %H:%M:%S").replace(tzinfo=CA.UTC).astimezone(CA.CT)
            term = WD.station_terminal(r["SITE_NO"])
            devs.append(abs(float(r["PPG"]) - rack[(term, t.date())] - 0.119))
    ok("SF10 every Riverbend diesel price is the TriState average for its pumping day plus the schedule's $0.119",
       max(devs) < 0.00006, "%d fills" % len(devs))
    aug = [v for (t_, d_), v in rack.items() if dt.date(2026, 8, 1) <= d_ <= dt.date(2026, 8, 27)]
    avg = sum(aug) / len(aug)
    gross_alt = sum(X["loads_acct"][a] * 8000 * (1.0055 * (DL.CHARGE[a] + avg) - (DL.ADDER + avg)) for a in DL.CHARGE)
    ok("SF15 the saving on the same number of gallons: reading the card side as the loads' gross gallons at the meter "
       "(0.55 per cent over net, at August's TriState average) moves the saving by $1,000 or more, so the prompt's "
       "'same number of gallons' closes it", abs(gross_alt - X["saving"]) > 1000 and "the same number of gallons" in
       ptxt, "%.0f at %.4f" % (gross_alt, avg))
    span_days = collections.Counter()
    for f in P.cards:
        if f["diesel"] and f["role"] == "sale" and f["unit"] in P.units and CA.BASE[0] <= f["t"].date() <= CA.BASE[1]:
            span_days[(f["acct"], WD.station_terminal(f["site"]) if f["site"] in WD.ST else "?")] += 1
    ok("SF16 the saving's price basis: the offer prices off the same TriState average the three card programmes bill "
       "from, so each gallon saves the card margin less $0.019 (the terminal and day spreads are not on the agreement "
       "terms)", "the same average your three card programmes bill from" in flat and
       X["per_gal"] == {"West": 0.1, "Central": 0.08, "East": 0.108}, str(dict(span_days)))
    tanks = openpyxl.load_workbook(os.path.join(target, R.F_TANKS), read_only=True).worksheets[0]
    usable = {row[0]: max(0, row[13] - row[6]) for row in list(tanks.iter_rows(values_only=True))[1:]}
    ok("SF12 the stock already in each tank (the 24 August dips over the unpumpable bottom) added to the draw: same call",
       max(usable.values()) <= 250 and counts({y: draw[y] + usable[y] for y in CA.YARDS}) == GOLD, str(usable))
    trow = {row[0]: row for row in list(tanks.iter_rows(values_only=True))[1:]}
    short = {y: max(0, trow[y][6] - trow[y][13]) for y in CA.YARDS}
    room_ok = all(trow[y][5] - max(trow[y][6], trow[y][13]) >= CA.LOAD + 2000 for y in CA.YARDS)
    alts = {"the old stock handed out first": {y: draw[y] - usable[y] for y in CA.YARDS},
            "a tank below its unpumpable bottom topped up from the first load": {y: draw[y] - short[y]
                                                                                 for y in CA.YARDS},
            "the heel kept under the last load": {y: draw[y] - trow[y][6] for y in CA.YARDS}}
    ok("SF14 the heel and the alarm: every tank takes a whole load over its unpumpable bottom with 2,000 gallons to "
       "spare under its overfill alarm, so room never binds; the dips' old stock handed out first, a bottom topped up "
       "from the first load, or the heel kept under the last load: same call",
       room_ok and all(counts(v) == GOLD for v in alts.values()) and
       len({(r[4], r[5], r[6]) for r in trow.values()}) >= 5,
       ", ".join("%s %s" % (k, tot(v)) for k, v in alts.items()))
    ok("SF13 the prompt marks only the 2025 programmes as last year's: Ask A's May, June and July are 2026's",
       ptxt.count("last year's") == 1 and "last year's Jackson and Knoxville" in ptxt and "May, June and July" in ptxt)
    ok("SF11 the as-of: the README's pull date and the offer's date agree (28 and 26 August 2026)",
       "August 28, 2026" in texts[DOC.F_README] and "August 26, 2026" in texts[DOC.F_OFFER])

    # ------------------------------------------------------------------ closure axes that can move
    swaps = [r for r in W.rename.values()]
    ok("AXIS 16 every device swap falls before the base (March and April 2026)",
       all(r[1] < CA.BASE[0] for r in swaps) and len(swaps) == 8)
    spares = [u for u, i in P.units.items() if i["status"] == "Spare"]
    spare_days = sum(len(P.days.get(u, {})) for u in spares)
    gone = {u.unit: u.sold for u in W.units if u.status == "Sold"}
    after = sum(1 for u, d0 in gone.items() for d in P.days.get(u, {}) if d >= d0)
    sold_geo = sum(len(P.stays.get(u, [])) for u in gone)
    ok("AXIS 19 the two spare trucks have no vehicle days, and the %d units sold in the autumn of 2024 none after "
       "their disposal and no yard visit in 2026: 'our own vehicles' has one reading" % len(gone),
       len(spares) == 2 and spare_days == 0 and len(gone) == 5 and after == 0 and sold_geo == 0 and
       all(P.units[u]["status"] == "Sold" for u in gone) and
       {i["status"] for i in P.units.values()} == {"Active", "Spare", "Sold"})
    open_ = 0
    for u, Sx in P.stays.items():
        for (a, b, g) in Sx:
            if a is not None and b is None and a.date() < dt.date(2026, 8, 20):
                open_ += 1
    ok("AXIS 15 every yard stay in the base is closed by an exit (no unpaired crossing)", open_ == 0 and P.unpaired == 0)
    sun = sum(1 for u in P.days if dt.date(2026, 8, 23) in P.days[u] or dt.date(2026, 6, 28) in P.days[u])
    ok("AXIS 12 no vehicle runs on the Sundays that bound the base: no run is open across either edge", sun == 0)
    ratio = {}
    for y in CA.YARDS:
        vals = [g for wk, g in weekly[y].items() if wk != dt.date(2026, 6, 29)]
        ratio[y] = max(vals) / min(vals)
    ok("AXIS 22 the base's regular weeks are steady at every yard (largest week within 15 per cent of the smallest)",
       max(ratio.values()) < 1.15, ", ".join("%s %.3f" % (CA.CODE[y], v) for y, v in ratio.items()))

    # ------------------------------------------------------------------ Ask A
    A = X["A"]
    gA = EXPECT["A"]
    ok("ASK A the nine account-months, by the local day each fill was pumped",
       all(round(A[k]) == v for k, v in gA.items()), str({k: round(v) for k, v in A.items()}))
    ok("ASK A every figure is 0.03 gallon or more from its rounding edge (each an exact sum of three-decimal gallons)",
       all(K.bin_clearance(v, 1) >= 0.03 for v in A.values()), str(["%.3f" % v for v in A.values()]))
    stops = {m: AS.ask_a(P, m) for m in ("asrecorded", "utc_only", "business_only", "allproducts")}
    natural = AS.ask_a(P, "asrecorded", "netted")
    netted = AS.ask_a(P, "golden", "netted")
    kept = AS.ask_a(P, "golden", "kept")
    miss = lambda st, keys: all(round(st[k]) != round(A[k]) for k in keys)
    hit = lambda st, keys: all(round(st[k]) == round(A[k]) for k in keys)
    allk = list(A)
    west = [k for k in allk if k[0] == "West"]
    rest = [k for k in allk if k[0] != "West"]
    # the truth: the world's diesel card fills by the local day they were pumped
    truth = collections.defaultdict(float)
    for u in W.units:
        for f in u.fills:
            if f["kind"] == "card":
                m_ = (f["t"].year, f["t"].month)
                truth[(R.ACCT_OF[u.yard], m_)] += round(f["gal"], 3)
    ok("ASK A the nine figures are the world's diesel bought by the local day it was pumped",
       all(abs(truth[k] - A[k]) < 0.01 for k in allk), str({k: round(truth[k]) for k in allk}))
    ok("ASK A each convention is filed in the programme's own document: Riverbend's UTC stamps (A5), its credits by "
       "transaction number (A6); Crossroads' Central-time business day closing at 10 pm (clause 5) and its reversals "
       "(clause 6); SMPM's business-day dates, site times and credit records (the layout)",
       all(x in flat for x in ("its date and time given in Coordinated Universal Time",
                               "the credit carries the transaction number of the duplicate it reverses",
                               "Dates and times on statements and in the portal are Central time",
                               "a purchase made at 10:00 pm or later carries the next day's date with the time it was "
                               "made", "carries the reference of the purchase it reverses",
                               "a sale made at 10:00 pm or later carries the next day's date",
                               "Transaction time at the member site, when the sale was made",
                               "A credit record reverses a detail record a member site sent twice")))
    # Crossroads states its clock (Central); four of its merchants sit on Eastern time
    eastern = {"CR42047", "CR42089", "CR42103", "CR42173"}
    alt_e = collections.defaultdict(float)
    n_e = 0
    for f in P.cards:
        if not f["diesel"] or f["role"] != "sale":
            continue
        d = AS.pump_date(f)
        if f["acct"] == "Central" and f["site"] in eastern:
            n_e += 1
            t_ct = dt.datetime.combine(d, f["t"].time())
            d = (t_ct + dt.timedelta(hours=1)).date()
        alt_e[(f["acct"], (d.year, d.month))] += f["gal"]
    ok("ASK A the %d Crossroads sales at merchants on Eastern time dated by the merchant's own clock: the same nine "
       "figures" % n_e, n_e > 500 and all(round(alt_e[k]) == round(A[k]) for k in A))
    late22 = sum(1 for f in P.cards if f["acct"] != "West" and f["t"].hour == 22 and f["t"].minute == 0)
    ok("ASK A the 10 pm boundary is filed as 'at 10:00 pm or later', so the %d sales stamped 22:00 have one reading"
       % late22, late22 > 0 and "after 10:00 pm" not in flat)
    ok("ASK A the natural read (every line summed on the date the files give it, DEF out) misses all nine",
       miss(natural, allk))
    ok("ASK A the dates as the files give them, the duplicates and credits handled, misses all nine",
       miss(stops["asrecorded"], allk))
    ok("ASK A every filed date convention applied but each credit netted on its own later date misses all nine (a "
       "duplicate day's gallons in one month, its credit in the next)", miss(netted, allk))
    dupm = sorted({(f["acct"], (AS.pump_date(f).year, AS.pump_date(f).month)) for f in P.cards
                   if f["role"] == "dup" and f["diesel"]} & set(allk))
    ok("ASK A the credits dropped and the duplicates kept misses each account-month with a duplicate day (%d)" %
       len(dupm), len(dupm) == 5 and miss(kept, dupm) and hit(kept, [k for k in allk if k not in dupm]),
       str(dupm))
    ok("ASK A Riverbend's UTC corrected but the 10 pm dates kept misses the six Crossroads and SMPM figures",
       miss(stops["utc_only"], rest) and hit(stops["utc_only"], west))
    ok("ASK A the 10 pm dates corrected but Riverbend's UTC read as local misses Delta's three",
       miss(stops["business_only"], west))
    ok("ASK A DEF counted as diesel misses all nine", miss(stops["allproducts"], allk))
    rs = AS.ask_a_resent(P)
    ok("ASK A the SMPM week received twice, kept twice, misses Holston's July by a week's gallons",
       rs[(2026, 7)] - A[("East", (2026, 7))] > 10000)
    before, after, unknown = AS.odometer_check(P)
    ok("ASK A fair: every Crossroads and SMPM fill timed 10 pm or later with a readable odometer was pumped the evening "
       "before its recorded date (odometer at or below that day's starting odometer)", before >= 1000 and after == 0,
       "%d before, %d after, %d unreadable" % (before, after, unknown))
    late_local = late_utc = n_w = 0
    for f in P.cards:
        if f["acct"] != "West" or not f["diesel"] or f["unit"] not in P.days:
            continue
        t_l = f["t"].replace(tzinfo=CA.UTC).astimezone(CA.CT)
        r_ = P.days[f["unit"]].get(t_l.date())
        if r_ is None:
            continue
        n_w += 1
        naive = f["t"].replace(tzinfo=CA.CT)
        late_local += naive > r_["b"] + dt.timedelta(hours=1)
        late_utc += t_l > r_["b"] + dt.timedelta(hours=1)
    ok("ASK A fair: read as local, a quarter or more of Riverbend's fills fall an hour or more after the vehicle's last "
       "ignition off that day; put on Central time, none do", late_local / n_w >= 0.25 and late_utc == 0,
       "%d of %d, %d" % (late_local, n_w, late_utc))

    # ------------------------------------------------------------------ Ask B
    Bb = X["B"]
    gB = EXPECT["B"]
    ok("ASK B the six programme months' net gallons and the dollars billed over the posted rack value of those gallons",
       all((round(Bb[k][0]), round(Bb[k][1])) == v for k, v in gB.items()),
       str({k: (round(v[0]), round(v[1])) for k, v in Bb.items()}))
    ok("ASK B every dollar figure is 3 cents or more from its rounding edge",
       all(K.bin_clearance(v[1], 1) >= 0.03 for v in Bb.values()), str([round(v[1], 3) for v in Bb.values()]))
    rack_o, drops, inv, cred = AS.read_overton(target)
    by = collections.defaultdict(list)
    for d in drops:
        by[d["bol"]].append(d)
    split_m = collections.Counter()
    straddle = 0
    for b_, v in by.items():
        if len(v) > 1:
            split_m[(AS.yard_of(v[0]["ship"]), (v[0]["delivered"].year, v[0]["delivered"].month))] += 1
            straddle += len({x["delivered"].month for x in v}) > 1
    ok("ASK B every programme month has a load dropped in two visits under one bill of lading, none across a month end",
       all(split_m[k] >= 1 for k in gB) and straddle == 0, str(dict(split_m)))
    bst = {m: AS.ask_b(target, m) for m in ("per_drop", "lift_date", "delivery_date", "no_credits", "by_lifting_month",
                                            "fuel_only", "no_rebills", "by_invoice_month")}
    ok("ASK B each convention is filed: the prompt's dollars billed over the posted rack value of the gallons after "
       "credits, taxes aside; the letters' posting at 6 pm Central with the next day's date, in effect when the load is "
       "lifted; the delivery export's times in Central (its headers); one invoice per bill of lading dated the day of "
       "delivery, its delivery service charges on it, a billing error credited in full and invoiced anew",
       "the dollars it billed over the posted rack value of those gallons after credits, taxes aside" in ptxt and
       all(x in flat for x in ("Overton posts its rack at 6:00 pm Central each day; each posting carries the next "
                               "day's date and prices every load lifted from then until 6:00 pm that day.",
                               "in effect when the load is lifted",
                               "One invoice for each bill of lading, dated the day of delivery: the load's net gallons "
                               "at the price above, any delivery service charges, and taxes and fees.",
                               "A billing error is corrected by a credit memo for the whole invoice and a new "
                               "invoice.")) and
       open(os.path.join(target, AS.F_DELIV), encoding="utf-8").readline().startswith(
           "Ticket,Terminal BOL,Terminal,Lifted (CT),Delivered (CT),"))
    ok("ASK B the delivery service charges (split drops, wait time) left out misses all six dollar figures",
       all(round(bst["fuel_only"][k][1]) != v[1] for k, v in gB.items()))
    ok("ASK B the credited invoices and their rebills both left out misses Knoxville July",
       round(bst["no_rebills"][("KNX", (2025, 7))][1]) != gB[("KNX", (2025, 7))][1])
    ok("ASK B dollars dated by invoice instead of by delivery: same twelve figures (each invoice dated its delivery, "
       "the rebills in the month of their loads)",
       all((round(bst["by_invoice_month"][k][0]), round(bst["by_invoice_month"][k][1])) == v for k, v in gB.items()))
    credited_inv = {c["ref"] for c in AS.read_overton(target)[3]}
    fr = [l for l in AS.read_overton(target)[2] if l["cls"] == "FREIGHT"]
    ok("ASK B every freight line is on an invoice credited in full: after credits no freight is billed", fr and
       all(l["inv"] in credited_inv for l in fr), "%d freight lines" % len(fr))
    ok("ASK B net summed over the delivery rows misses all six gallon figures and all six dollar figures",
       all(round(bst["per_drop"][k][0]) != v[0] and round(bst["per_drop"][k][1]) != v[1] for k, v in gB.items()))
    m_lift = [k for k, v in gB.items() if round(bst["lift_date"][k][1]) != v[1]]
    m_del = [k for k, v in gB.items() if round(bst["delivery_date"][k][1]) != v[1]]
    ok("ASK B the posting dated the lifting day and the posting dated the delivery day each miss four or more of the "
       "six dollar figures, and every month is missed by one of them",
       len(m_lift) >= 4 and len(m_del) >= 4 and set(m_lift) | set(m_del) == set(gB), "%d, %d" % (len(m_lift), len(m_del)))
    ok("ASK B the July credit memos unread misses Knoxville July by the two rebilled loads",
       round(bst["no_credits"][("KNX", (2025, 7))][1]) - gB[("KNX", (2025, 7))][1] > 30000)
    ok("ASK B loads dated by lifting move the load lifted on 30 June into June",
       round(bst["by_lifting_month"][("KNX", (2025, 6))][0]) != gB[("KNX", (2025, 6))][0])
    fair = []
    promo = {"JXN": 0.024, "KNX": 0.016}
    first = {}
    for d in sorted(drops, key=lambda x: x["delivered"]):
        first.setdefault(d["bol"], d)
    credited = {c["ref"] for c in cred}
    for line in inv:
        if line["cls"] != "PRODUCT" or line["inv"] in credited:
            continue
        d = first[line["bol"]]
        eff = d["lifted"].date() + dt.timedelta(days=1) if d["lifted"].hour >= 18 else d["lifted"].date()
        fair.append(abs(line["amt"] / line["qty"] - rack_o[(d["terminal"], eff)] - promo[AS.yard_of(d["ship"])]))
    ok("ASK B fair: every billed fuel price is Overton's posting in effect at lifting (6 pm Central for the next day) "
       "plus the programme margin", max(fair) < 0.0001, "%d loads" % len(fair))
    # the cracker: the call landed, every filed convention applied, each line summed where the files date it
    crack_b = bst["per_drop"]
    ok("CRACKER the careful route (stack, dedupe the resent week, DEF out, every filed date convention applied, each "
       "credit summed on its own date; delivery rows as listed) misses every ask figure",
       miss(netted, allk) and all(round(crack_b[k][0]) != v[0] and round(crack_b[k][1]) != v[1]
                                  for k, v in gB.items()))

    # ------------------------------------------------------------------ separation
    ok("SEPARATION the main call's rows carry no device: the golden reads no card export, no Overton file and no rack",
       counts(CA.draw_by_yard(P_nocard)) == GOLD)

    # ------------------------------------------------------------------ realism sweep
    figs = [round(v) for v in A.values()] + [round(v[0]) for v in Bb.values()] + [round(v[1]) for v in Bb.values()] + \
        [round(v) for v in draw.values()]
    ok("REALISM no two graded figures or yard draws are equal", len(figs) == len(set(figs)), figs)
    ok("REALISM no yard draw or account-month is a round hundred", all(f % 100 for f in [round(v) for v in draw.values()]
                                                                      + [round(v) for v in A.values()]))
    ifta = openpyxl.load_workbook(os.path.join(target, R.F_IFTA), read_only=True).worksheets[0]
    mpg = collections.defaultdict(set)
    for row in list(ifta.iter_rows(values_only=True))[1:]:
        mpg[row[0]].add((row[2], row[6]))
    ok("REALISM the IFTA returns' fleet MPG differs between the licences in each quarter",
       all(len({m for _, m in v}) == len({a for a, _ in v}) for v in mpg.values()))
    ins_bad = [u.unit for u in W.units if not K.working(u.in_service) or
               (P.days.get(u.unit) and min(P.days[u.unit]) < u.in_service)]
    ok("REALISM every unit's in-service date is a working day before its first vehicle day", not ins_bad, ins_bad)
    names_ = set(K.PEOPLE.values())
    firsts = {n.split()[0] for n in names_}
    found = set(re.findall(r"\b([A-Z][a-z]+ [A-Z][a-z]+)\b", alltext))
    persons = {n for n in found if n.split()[0] in firsts}
    ok("REALISM every person named in the documents is a drawn persona", persons <= names_, sorted(persons - names_))
    dash = [f for f in os.listdir(target) if f.endswith((".csv", ".md", ".txt")) and (
        "\u2014" in open(os.path.join(target, f), encoding="utf-8").read())]
    dash += [f for f, t in texts.items() if "\u2014" in t or "\u2013" in t]
    sub = open(os.path.join(os.path.dirname(os.path.abspath(target)), "submission.md"), encoding="utf-8").read()
    ok("REALISM no em or en dash in the pack or the submission", not dash and "\u2014" not in sub and "\u2013" not in sub,
       dash)
    stamp = K.EXPORT_STAMP.timestamp()
    ok("REALISM every shipped file carries the pull's timestamp", all(
        abs(os.path.getmtime(os.path.join(target, f)) - stamp) < 1 for f in os.listdir(target)))
    late = [f for f, (c, m) in DOC.PDF_STAMPS.items() if m[2:10].decode() > "20260828"]
    for f in (R.F_UNITS, R.F_CENTRAL, R.F_TANKS, R.F_IFTA):
        with __import__("zipfile").ZipFile(os.path.join(target, f)) as z:
            core = z.read("docProps/core.xml").decode()
        if re.search(r"<dcterms:modified[^>]*>([^<]+)<", core).group(1)[:10] > "2026-08-28":
            late.append(f)
    ok("REALISM no document or workbook is dated after the pull", not late, late)
    readme = texts[DOC.F_README]
    unlisted = [f for f in sorted(os.listdir(target)) if f != DOC.F_README and f not in readme]
    ok("REGISTER the README names every shipped file", not unlisted, unlisted)
    tz_words = re.findall(r"\bUTC\b|time zone|Central time|Eastern time|business day|10 pm", readme + texts[DOC.F_NOTES])
    ok("REGISTER the README and the yard notes narrate no clock or business-day convention", not tz_words, tz_words)

    # ------------------------------------------------------------------ input gates and distractors
    files = sorted(os.listdir(target))
    fmts = {f.rsplit(".", 1)[1] for f in files}
    big = max(B.counts.values())
    ok("GATE 10 or more files, 3 or more formats, one of 25,000 or more rows",
       len(files) >= 10 and len(fmts) >= 3 and big >= 25000, "%d files, %s, %d rows" % (len(files), sorted(fmts), big))
    meta = json.load(open(meta_path))
    ok("GATE two distractors named in metadata.json and shipped, neither named as one",
       len(meta["distractor_files"]) >= 2 and all(d in files for d in meta["distractor_files"]) and
       not any("distract" in f for f in files))
    tmp = tempfile.mkdtemp(dir=os.path.dirname(os.path.abspath(target)), suffix=".tmp")
    try:
        for f in files:
            if f not in meta["distractor_files"]:
                os.symlink(os.path.abspath(os.path.join(target, f)), os.path.join(tmp, f))
        P4 = CA.read(tmp)
        G4 = CA.golden(P4)
        A4 = AS.ask_a(P4)
        B4 = AS.ask_b(tmp)
        ok("GATE the golden and both asks unchanged with the distractors deleted",
           G4["loads"] == loads and all(abs(A4[k] - A[k]) < 1e-6 for k in A) and
           all(abs(B4[k][0] - Bb[k][0]) < 1e-6 and abs(B4[k][1] - Bb[k][1]) < 1e-6 for k in Bb))
    finally:
        shutil.rmtree(tmp)

    # ------------------------------------------------------------------ the goldens agree with figures.json
    F = json.load(open(os.path.join(os.path.dirname(os.path.abspath(target)), "generator", "figures.json")))
    wb = openpyxl.load_workbook(os.path.join(golden, DL.XLSX), data_only=True)
    ws = wb["Loads"]
    rows = {ws.cell(row=i, column=1).value: [ws.cell(row=i, column=c).value for c in range(1, 10)] for i in range(6, 13)}
    ok("GOLDEN the workbook's counts, total and saving are the figures'",
       all(rows[y][3] == F["call"]["loads_by_yard"][y] for y in CA.YARDS) and rows["Total"][3] == F["call"]["total_loads"]
       and "{:,}".format(rows["Total"][8]) == F["call"]["saving_usd"])
    wa = wb["Card gallons"]
    vals = [[wa.cell(row=5 + i, column=2 + j).value for j in range(3)] for i in range(3)]
    ok("GOLDEN the workbook's card gallons are the figures'", all(
        "{:,}".format(vals[i][j]) == F["answers"]["ask_a"]["%s 2026-%02d" % (a, m)]
        for i, a in enumerate(("West", "Central", "East")) for j, m in enumerate((5, 6, 7))))
    wbb = wb["2025 programmes"]
    okb = True
    for i in range(6):
        key = ["JXN 2025-03", "JXN 2025-04", "JXN 2025-05", "KNX 2025-06", "KNX 2025-07", "KNX 2025-08"][i]
        okb &= "{:,}".format(wbb.cell(row=5 + i, column=4).value) == F["answers"]["ask_b"][key]["net_gal"]
        okb &= "{:,}".format(wbb.cell(row=5 + i, column=5).value) == F["answers"]["ask_b"][key]["usd"]
    ok("GOLDEN the workbook's programme months are the figures'", okb)
    ok("GOLDEN the chart is rendered", os.path.getsize(os.path.join(golden, DL.PNG)) > 20000)

    fails = [x for x in RESULTS if not x[1]]
    for label, good, detail in RESULTS:
        print("%s  %s%s" % ("PASS" if good else "FAIL", label, ("  [%s]" % detail) if detail else ""))
    print("checks: %d assertions, %d failed" % (len(RESULTS), len(fails)))
    if fails:
        raise SystemExit("checks failed: %s" % "; ".join(x[0] for x in fails))
    return RESULTS
