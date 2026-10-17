"""The goldens, from the shipped files only: overton_loads_sep-nov_2026.xlsx, yard_island_draw.png, submission.md and
generator/figures.json (every figure the write-up states, as it states it)."""
from __future__ import annotations

import collections
import datetime as dt
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import asks as A
import calc as C
import postprocess as PP

XLSX = "overton_loads_sep-nov_2026.xlsx"
PNG = "yard_island_draw.png"
ACCT_NAME = {"West": "Delta, Riverbend direct bill", "Central": "Keeble, Crossroads Fleet Card",
             "East": "Holston, SMPM fleet card"}
ACCT_SHORT = {"West": "Riverbend", "Central": "Crossroads", "East": "SMPM"}
CHARGE = {"West": 0.119, "Central": 0.084 + 0.015, "East": 0.127}     # the schedules' charge over the average
ADDER = 0.019
MONTH = {5: "May", 6: "June", 7: "July", 3: "March", 4: "April", 8: "August"}
PROG = {"JXN": "Jackson", "KNX": "Knoxville"}


def fmt(x, dp=0):
    return format(round(x, dp) if dp else int(round(x)), ",.%df" % dp if dp else ",")


def compute(target):
    P = C.read(target)
    G = C.golden(P)
    R, S = C.rungs(P, G["draw"])
    out = dict(P=P, G=G, R=R, S=S)
    out["A"] = A.ask_a(P)
    out["B"] = A.ask_b(target)
    # the build behind each yard's draw
    tz_units = collections.defaultdict(list)
    for u, info in P.units.items():
        if info["status"] == "Active":
            tz_units[info["yard"]].append(u)
    cap_r = G["reserve"]
    build = {}
    for y in C.YARDS:
        visits = fuel = fills = 0.0
        for u in tz_units[y]:
            cap = P.units[u]["tank"] * (1.0 - cap_r)
            for (a, g) in C.island_fills(P, u, cap):
                if C.BASE[0] <= a.date() <= C.BASE[1]:
                    visits += 1
                    fills += g
            fuel += sum(r["fuel"] for d, r in P.days.get(u, {}).items() if C.BASE[0] <= d <= C.BASE[1])
        build[y] = dict(units=len(tz_units[y]), visits=int(visits), fuel=fuel, fills=fills, draw=G["draw"][y],
                        loads=G["loads"][y], card=R["R0 card gallons, trailing eight weeks"][y])
    out["build"] = build
    # the findings the call rests on: Nashville's sleepers in the yard at midday every Wednesday; Blountville's
    # sleepers and Cookeville's vans in the yard once a week
    mid = {}
    for u in tz_units["Nashville"]:
        if P.units[u]["model"] != "T68SL":
            continue
        n = 0
        for (s0, s1, gname) in P.stays.get(u, []):
            if gname == "KSC Nashville Yard" and s0 and s1:
                a = s0.astimezone(C.CT)
                b = s1.astimezone(C.CT)
                if a.date() == b.date() and C.BASE[0] <= a.date() <= C.BASE[1] and a.weekday() == 2:
                    n += 1
        mid[u] = n
    out["nsh_sleepers"] = len(mid)
    out["nsh_midday_weeks"] = min(mid.values())
    weekly = {}
    for y, model in (("Blountville", "T68SL"), ("Cookeville", "SPR25")):
        per = []
        for u in tz_units[y]:
            if P.units[u]["model"] != model:
                continue
            per.append(sum(1 for (a, g) in C.island_fills(P, u, 1e9) if C.BASE[0] <= a.date() <= C.BASE[1]))
        weekly[y] = (len(per), min(per), max(per))
    out["weekly"] = weekly
    # the saving
    loads_acct = collections.defaultdict(int)
    for y in C.YARDS:
        loads_acct[C.ACCOUNT[y]] += G["loads"][y]
    per_load = {a: round(C.LOAD * (CHARGE[a] - ADDER), 2) for a in CHARGE}
    out["per_load"] = per_load
    out["per_gal"] = {a: round(CHARGE[a] - ADDER, 3) for a in CHARGE}
    out["saving_acct"] = {a: loads_acct[a] * per_load[a] for a in CHARGE}
    out["saving"] = sum(out["saving_acct"].values())
    out["loads_acct"] = dict(loads_acct)
    return out


# ------------------------------------------------------------------------------------------------- the workbook
THIN = Side(style="thin", color="A3ABB5")
HEAD = PatternFill("solid", fgColor="DCE3EA")
TOT = PatternFill("solid", fgColor="F1F4F7")


def _head(ws, row, cols, widths=None):
    for c, v in enumerate(cols, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = Font(bold=True)
        cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = Border(bottom=THIN)
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w


def write_xlsx(path, X):
    G, B, Aa = X["G"], X["B"], X["A"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Loads"
    ws["A1"] = "Contract %d Overton loads for the yard islands, 14 September to 8 November 2026" % G["total"]
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Count to confirm to Overton by Friday 4 September. Each yard takes the whole 8,000-gallon loads its "
                "island will hand out to our own vehicles by 8 November; the saving is against the same gallons on "
                "the card accounts.")
    ws["A3"] = "Data through Thursday 27 August 2026. Definitions and sources on the Notes sheet."
    ws["A3"].font = Font(italic=True, color="5F6670")
    _head(ws, 5, ["Yard", "Card account", "Island draw, 14 Sep to 8 Nov (gal)", "Loads", "Gallons in the loads",
                  "Card charge over TriState average ($/gal)", "Overton over TriState average ($/gal)",
                  "Saving ($/gal)", "Saving ($)"], [13, 30, 16, 8, 13, 16, 16, 10, 12])
    r = 6
    for y in C.YARDS:
        a = C.ACCOUNT[y]
        n = G["loads"][y]
        vals = [y, ACCT_NAME[a], round(G["draw"][y]), n, n * 8000, CHARGE[a], ADDER, X["per_gal"][a],
                round(n * X["per_load"][a])]
        for c, v in enumerate(vals, 1):
            ws.cell(row=r, column=c, value=v)
        r += 1
    tot = ["Total", "", round(sum(G["draw"].values())), G["total"], G["total"] * 8000, None, None, None,
           round(X["saving"])]
    for c, v in enumerate(tot, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(bold=True)
        cell.fill = TOT
        cell.border = Border(top=THIN)
    for rr in range(6, r + 1):
        ws.cell(row=rr, column=3).number_format = "#,##0"
        ws.cell(row=rr, column=5).number_format = "#,##0"
        for c in (6, 7, 8):
            ws.cell(row=rr, column=c).number_format = "0.000"
        ws.cell(row=rr, column=9).number_format = "$#,##0"
    ws.freeze_panes = "A6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = "A1:I%d" % r

    w2 = wb.create_sheet("Card gallons")
    w2["A1"] = "Diesel gallons bought on the card accounts, by the local day of the fill"
    w2["A1"].font = Font(bold=True, size=12)
    w2["A2"] = ("Riverbend's file times are UTC and are put on Central time; Crossroads and SMPM date a truck-stop fill "
                "after 10 pm with the next business day, so those fills are put back on the evening they were pumped. "
                "DEF is excluded; the SMPM batch for 6 to 12 July received twice is counted once.")
    w2["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    w2.row_dimensions[2].height = 45
    _head(w2, 4, ["Card account", "May 2026", "June 2026", "July 2026"], [34, 13, 13, 13])
    for k, a in enumerate(("West", "Central", "East")):
        w2.cell(row=5 + k, column=1, value=ACCT_NAME[a])
        for j, m in enumerate(A.MONTHS):
            c = w2.cell(row=5 + k, column=2 + j, value=round(Aa[(a, m)]))
            c.number_format = "#,##0"
    w2.merge_cells("A2:D2")

    w3 = wb.create_sheet("2025 programmes")
    w3["A1"] = "Overton's 2025 island programmes: net gallons delivered and billed above its posted rack, by month"
    w3["A1"].font = Font(bold=True, size=12)
    w3["A2"] = ("Each load once, in the month of its first drop. Billed above the rack: the invoices' product, freight "
                "and service lines, net of the July credit memos, less each load's net gallons at Overton's posting "
                "in effect when it was lifted (posted 6 pm Central for the next day). Taxes aside.")
    w3["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    w3.row_dimensions[2].height = 45
    _head(w3, 4, ["Programme", "Month", "Loads", "Net gallons delivered", "Billed above posted rack ($)"],
          [13, 11, 8, 15, 16])
    rr = 5
    for y in ("JXN", "KNX"):
        for m in A.PROG_MONTHS[y]:
            g, d = B[(y, m)]
            vals = [PROG[y], "%s %d" % (MONTH[m[1]], m[0]), X["B_loads"][(y, m)], round(g), round(d)]
            for c, v in enumerate(vals, 1):
                w3.cell(row=rr, column=c, value=v)
            w3.cell(row=rr, column=4).number_format = "#,##0"
            w3.cell(row=rr, column=5).number_format = "$#,##0"
            rr += 1
    w3.merge_cells("A2:E2")

    w4 = wb.create_sheet("Draw build")
    w4["A1"] = "Island draw by yard"
    w4["A1"].font = Font(bold=True, size=12)
    w4["A2"] = ("At each yard visit the island puts aboard the fuel the vehicle burned since its last visit, up to its "
                "tank less the %.1f%% reserve the 2025 islands show sleepers coming back on (%d fills over 225 gallons, "
                "average %.1f of 300). Base: the eight weeks 29 June to 23 August (39 working days), scaled to the "
                "window's 40." % (100 * G["reserve"], G["n_reserve"], 300 * (1 - G["reserve"])))
    w4["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    w4.row_dimensions[2].height = 45
    _head(w4, 4, ["Yard", "Units", "Yard visits, 29 Jun to 23 Aug", "Fuel burned, 29 Jun to 23 Aug (gal)",
                  "Island fills, 29 Jun to 23 Aug (gal)", "Island draw, 14 Sep to 8 Nov (gal)", "Loads",
                  "Card gallons, 29 Jun to 23 Aug"], [13, 7, 13, 15, 15, 15, 7, 14])
    rr = 5
    for y in C.YARDS:
        b = X["build"][y]
        vals = [y, b["units"], b["visits"], round(b["fuel"]), round(b["fills"]), round(b["draw"]), b["loads"],
                round(b["card"])]
        for c, v in enumerate(vals, 1):
            cell = w4.cell(row=rr, column=c, value=v)
            if c >= 3 and c != 7:
                cell.number_format = "#,##0"
        rr += 1
    w4.merge_cells("A2:H2")

    w5 = wb.create_sheet("Notes")
    w5.column_dimensions["A"].width = 22
    w5.column_dimensions["B"].width = 110
    notes = [
        ("As of", "Telematics and card data through Thursday 27 August 2026, pulled Friday 28 August."),
        ("Load", "8,000 net gallons, lifted for and delivered to one yard tank (Overton's offer of 26 August), so "
                 "whole loads are counted yard by yard."),
        ("Island draw", "At each yard visit (the telematics geofence export) the island puts aboard the fuel the "
                        "vehicle burned since its previous visit (Fuel Used in the vehicle-day export, a day split by "
                        "driving time either side of a midday visit), up to its tank on the unit list's model row "
                        "less the reserve drivers come back on."),
        ("Reserve", "%.1f per cent of the tank: the %d fills over 225 gallons in the 2025 island transactions, every "
                    "one a T680 coming back off the road, averaged %.1f of 300 gallons." % (
                        100 * G["reserve"], G["n_reserve"], 300 * (1 - G["reserve"]))),
        ("Base and window", "The eight weeks 29 June to 23 August 2026 (39 working days, Friday 3 July off) scaled "
                            "to the window's 40 working days. No route, fleet or yard change is planned before "
                            "December (yard notes)."),
        ("Saving", "Per Overton's offer, each gallon saves the card programme's charge over the TriState average "
                   "less Overton's 1.9 cents: Riverbend $0.119; Crossroads $0.084 and its $0.015 network fee; SMPM "
                   "$0.127 (the programmes' pricing schedules). Fuel taxes are paid either way and are left out."),
        ("Card gallons", "Diesel only, DEF left out. Riverbend's download stamps fills in UTC, put on Central time. "
                         "Crossroads and SMPM give a truck-stop fill after 10 pm the next business day; every such "
                         "fill with a readable odometer shows one at or below that day's starting odometer, so these "
                         "fills are put back on the evening they were pumped. SMPM's batch for 6 to 12 July came in "
                         "twice and is counted once."),
        ("2025 programmes", "Each terminal bill of lading once, in the month of its first drop: a load dropped in "
                            "two visits shows the bill's net gallons on both delivery rows. The two July Knoxville "
                            "invoices credited on 21 July (a freight-zone charge billed in error) are replaced by "
                            "their rebills. Above the rack is against Overton's posting in effect at lifting, posted "
                            "at 6 pm Central for the next day."),
    ]
    _head(w5, 1, ["Item", "Note"])
    for i, (k, v) in enumerate(notes, 2):
        w5.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = w5.cell(row=i, column=2, value=v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        w5.cell(row=i, column=1).alignment = Alignment(vertical="top")
        w5.row_dimensions[i].height = 15 * (1 + len(v) // 105)
    wb.save(path)
    PP.normalize_ooxml(path, "Fleet Cost", "Fleet Cost", "2026-08-31T14:05:00Z", "2026-09-02T20:41:00Z",
                       excel=True, company="Keeble Supply Co.")


# ------------------------------------------------------------------------------------------------- the chart
def write_png(path, X):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    for f in ("/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf",
              "/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf"):
        font_manager.fontManager.addfont(f)
    plt.rcParams["font.family"] = "Carlito"
    G = X["G"]
    navy, pale, grey = "#1f4e79", "#c3cfdc", "#5f6670"
    fig = plt.figure(figsize=(10.0, 6.0), dpi=120)
    ax = fig.add_axes([0.075, 0.17, 0.9, 0.62])
    xs = list(range(len(C.YARDS)))
    w = 0.36
    draw = [G["draw"][y] / 1000 for y in C.YARDS]
    card = [X["build"][y]["card"] / 1000 for y in C.YARDS]
    ax.bar([x - w / 2 for x in xs], draw, w, color=navy, label="Island draw, 14 Sep to 8 Nov 2026 (forecast)")
    ax.bar([x + w / 2 for x in xs], card, w, color=pale, label="Bought on the cards, 29 Jun to 23 Aug 2026")
    for x, y in zip(xs, C.YARDS):
        n = G["loads"][y]
        for k in range(1, n + 1):
            ax.plot([x - w + 0.025, x - 0.025], [8 * k, 8 * k], color="white", lw=1.2, solid_capstyle="butt")
        ax.text(x - w / 2, G["draw"][y] / 1000 + 1.0, "%d load%s" % (n, "" if n == 1 else "s"), ha="center",
                va="bottom", fontsize=10, color=navy, fontweight="bold")
    notes = {"Nashville": "back midweek\nfor a second fill", "Cookeville": "lot vans: one\ntankful a week",
             "Blountville": "4-night loops: one\ntankful a week"}
    for x, y in zip(xs, C.YARDS):
        if y in notes:
            top = max(G["draw"][y] / 1000 + 4.4, X["build"][y]["card"] / 1000 + 1.0)
            ax.text(x + w / 2, top, notes[y], ha="center", va="bottom", fontsize=8.6, color=grey)
    ax.set_xticks(xs)
    ax.set_xticklabels(["%s\n%s" % (y, ACCT_SHORT[C.ACCOUNT[y]]) for y in C.YARDS], fontsize=10)
    ax.set_ylabel("Thousand gallons", fontsize=10)
    ax.set_ylim(0, 80)
    ax.set_yticks(range(0, 81, 10))
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    ax.grid(axis="y", color="#e6e9ed", lw=0.8)
    ax.set_axisbelow(True)
    ax.legend(frameon=False, fontsize=9.5, loc="upper left", ncol=2, bbox_to_anchor=(0.0, 1.07))
    fig.text(0.075, 0.935, "The yard islands will hand out %d whole Overton loads by 8 November" % G["total"],
             fontsize=15, fontweight="bold", color="#1f2329")
    fig.text(0.075, 0.895, "Each island fuels a vehicle only when it is in the yard, at most a tankful a visit. White "
             "ticks mark each whole 8,000-gallon load inside a yard's draw.", fontsize=10, color=grey)
    fig.text(0.075, 0.035, "Source: telematics vehicle days and yard geofence visits, 29 June to 23 August 2026, "
             "scaled from 39 to 40 working days; Riverbend, Crossroads and SMPM diesel, DEF excluded.",
             fontsize=8.6, color=grey)
    fig.savefig(path, dpi=120, metadata={"Software": None})
    plt.close(fig)


# ------------------------------------------------------------------------------------------------- the write-up
def submission(X):
    G, R, B, Aa = X["G"], X["R"], X["B"], X["A"]
    L = G["loads"]
    yard_list = ", ".join("%s %d" % (y, L[y]) for y in C.YARDS)
    r0 = sum(C.floor_loads(v) for v in R["R0 card gallons, trailing eight weeks"].values())
    r1 = sum(C.floor_loads(v) for v in R["R1a pooled 2025 share"].values())
    r2 = sum(C.floor_loads(v) for v in R["R2a tankful per run between nights at the yard"].values())
    X["rivals"] = dict(card_gallons=r0, share_2025=r1, overnight_cut=r2)
    s = []
    s += ["## Tags", "",
          "**Domain:** General Statistics & Data Science (resampling and simulation methods: a fleet's yard-island fuel "
          "draw replayed from each vehicle's recent duty, stretch by stretch between yard visits, each stretch capped "
          "at a tankful).",
          "**Analytical objective:** Forecasting & Predictive Modeling.", ""]
    s += ["## 1. Final Recommendation", "",
          "**Contract %d of Overton's discounted loads for 14 September to 8 November (%s), saving $%s against the "
          "same gallons on the card accounts.**" % (G["total"], yard_list, fmt(X["saving"])),
          "An island fuels a vehicle only at its yard visits, at most the tankful above the reserve its driver keeps, "
          "so each yard's draw is built from the fuel each vehicle burns between visits. Not the %d loads the "
          "vehicles' card gallons fill, which count the fuel Blountville's four-night sleepers and Cookeville's "
          "lot-based vans have to buy on the road; not %d from the 2025 islands' share of the same vehicles' fuel, "
          "measured on Jackson and Knoxville fleets mostly back every day or between one-night loops; and not %d from "
          "capping each run between nights at the yard, which misses Nashville's midweek reload." % (r0, r1, r2), ""]
    wk = X["weekly"]
    s += ["## 2. Critical Components", "",
          "1. Nashville's **%d** sleepers are in the yard at midday every Wednesday, so each takes two island fills a "
          "week and the island keeps nearly all of their fuel." % X["nsh_sleepers"],
          "2. Blountville's **%d** sleepers on four-night loops and Cookeville's **%d** vans based at the Sparta and "
          "Crossville lots are in the yard once a week, so each takes at most one tankful a week." % (
              wk["Blountville"][0], wk["Cookeville"][0]),
          "3. Each load saves **$%s** on Delta's Riverbend account (Memphis, Jackson), **$%s** on Crossroads "
          "(Nashville, Cookeville, its network fee included) and **$%s** on SMPM (Knoxville, Blountville)." % (
              fmt(X["per_load"]["West"]), fmt(X["per_load"]["Central"]), fmt(X["per_load"]["East"])), ""]
    draws = ", ".join("%s %s" % (y, fmt(G["draw"][y])) for y in C.YARDS)
    s += ["## 3. Step-by-Step Solution", "",
          "1. Took the %d active units, their yards and their models' fuel tanks from fleet_units_2026-08.xlsx, and "
          "named each telematics row by its vehicle, through telematics_device_assignments.csv where a new device "
          "reported under its serial." % sum(b["units"] for b in X["build"].values()),
          "2. Paired each unit's Enter and Exit rows at its own yard's geofence in "
          "geofence_events_2026-03-01_2026-08-27.csv into yard visits, and cut its days in "
          "vehicle_days_2024-09-01_2026-08-27.csv at them, a day's Fuel Used split by its driving time either side of "
          "a midday visit.",
          "3. In fms_island_transactions_2025.csv (each resent row once), the %d fills over 225 gallons, every one a "
          "T680 sleeper coming back off the road, averaged %.1f of 300 gallons: each vehicle's fill is capped at its "
          "tank less %.1f%%." % (G["n_reserve"], 300 * (1 - G["reserve"]), 100 * G["reserve"]),
          "4. For every yard visit from 29 June to 23 August, took the smaller of the fuel burned since the unit's "
          "previous yard visit and its capped tankful, summed by yard and scaled from those weeks' 39 working days "
          "to the window's 40: %s gallons." % draws,
          "5. Took each yard's whole 8,000-gallon loads: %s, %d in all." % (yard_list, G["total"]),
          "6. Priced each load at its account's charge over the TriState average in "
          "card_programmes_pricing_schedules.pdf less Overton's $0.019 in overton_offer_2026-08-26.pdf, on its 8,000 "
          "gallons: %d loads save $%s, so contract %d." % (G["total"], fmt(X["saving"]), G["total"]), ""]
    s += ["## 4. Deliverable Answers", "", "### %s" % XLSX, "",
          "1. Loads to contract for 14 September to 8 November, by yard: %s." % yard_list,
          "2. Total: %d loads." % G["total"],
          "3. Saving against the same gallons on the card accounts at their agreement terms: $%s." % fmt(X["saving"]),
          "4. Diesel gallons each card account's vehicles bought, by month:"]
    for a in ("West", "Central", "East"):
        s.append("   - %s: May %s gallons, June %s gallons, July %s gallons" % (
            ACCT_NAME[a], fmt(Aa[(a, (2026, 5))]), fmt(Aa[(a, (2026, 6))]), fmt(Aa[(a, (2026, 7))])))
    s.append("5. Overton's 2025 island programmes, net gallons delivered and dollars billed above its posted rack, "
             "taxes aside:")
    for y in ("JXN", "KNX"):
        for m in A.PROG_MONTHS[y]:
            g, d = B[(y, m)]
            s.append("   - %s, %s %d: %s gallons, $%s" % (PROG[y], MONTH[m[1]], m[0], fmt(g), fmt(d)))
    s += ["", "### %s" % PNG, "",
          "6. Bar chart: each yard's island draw for 14 September to 8 November beside its vehicles' card gallons for "
          "29 June to 23 August, white ticks marking each whole 8,000-gallon load inside the draw bar and each "
          "yard's load count labelled above it.", ""]
    return "\n".join(s)


def figures(X):
    G, B, Aa = X["G"], X["B"], X["A"]
    F = {
        "call": {"total_loads": G["total"], "loads_by_yard": {y: G["loads"][y] for y in C.YARDS},
                 "saving_usd": fmt(X["saving"])},
        "components": {"nashville_sleepers": X["nsh_sleepers"], "nashville_midday_weeks": X["nsh_midday_weeks"],
                       "blountville_sleepers": X["weekly"]["Blountville"][0],
                       "cookeville_vans": X["weekly"]["Cookeville"][0],
                       "saving_per_load_usd": {a: fmt(v) for a, v in X["per_load"].items()},
                       "saving_per_gallon": X["per_gal"]},
        "steps": {"active_units": sum(b["units"] for b in X["build"].values()),
                  "reserve_fills": G["n_reserve"], "reserve_mean_fill_gal": "%.1f" % (300 * (1 - G["reserve"])),
                  "reserve_pct": "%.1f" % (100 * G["reserve"]), "fill_threshold_gal": 225, "tank_gal": 300,
                  "base_working_days": 39, "window_working_days": 40, "load_net_gal": 8000, "adder": "0.019",
                  "draw_gal": {y: fmt(G["draw"][y]) for y in C.YARDS},
                  "build": {y: {k: (round(v) if isinstance(v, float) else v) for k, v in b.items()}
                            for y, b in X["build"].items()}},
        "rivals": X["rivals"],
        "answers": {"ask_a": {"%s %d-%02d" % (a, m[0], m[1]): fmt(v) for (a, m), v in Aa.items()},
                    "ask_b": {"%s %d-%02d" % (y, m[0], m[1]): {"net_gal": fmt(g), "usd": fmt(d)}
                              for (y, m), (g, d) in B.items()}},
        "dates": {"base": "29 June to 23 August", "window": "14 September to 8 November"},
    }
    return F


def main(target, golden, fig_path, sub_path, B=None):
    os.makedirs(golden, exist_ok=True)
    X = compute(target)
    rack, drops, inv, cred = A.read_overton(target)
    first = {}
    for d in sorted(drops, key=lambda x: x["delivered"]):
        first.setdefault(d["bol"], d)
    X["B_loads"] = collections.Counter((A.yard_of(d["ship"]), (d["delivered"].year, d["delivered"].month))
                                       for d in first.values())
    write_xlsx(os.path.join(golden, XLSX), X)
    write_png(os.path.join(golden, PNG), X)
    text = submission(X)
    with open(sub_path, "w", encoding="utf-8") as fh:
        fh.write(text)
    with open(fig_path, "w", encoding="utf-8") as fh:
        json.dump(figures(X), fh, indent=2)
        fh.write("\n")
    return X
