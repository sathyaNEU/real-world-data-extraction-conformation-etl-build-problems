"""The call and every rung, computed from the shipped files only (the solver's side of the build).

The golden: each yard island hands each vehicle, at each of its yard visits, the smaller of the fuel it burned since its
last yard visit and the tankful above the reserve its driver keeps. The stretches are built from the telematics vehicle
days (fuel, engine-on interval) and the geofence export's yard entries and exits; the tank is the model row's; the
reserve is what the 2025 islands show long-run sleepers coming back on. The base is the eight complete weeks to Sunday
23 August, scaled from its 39 working days to the window's 40; each yard takes the whole loads of its draw."""
from __future__ import annotations

import collections
import csv
import datetime as dt
import math
import os
import re
from zoneinfo import ZoneInfo

import openpyxl

CT, ET, UTC = ZoneInfo("America/Chicago"), ZoneInfo("America/New_York"), dt.timezone.utc
F_DAYS = "vehicle_days_2024-09-01_2026-08-27.csv"
F_GEO = "geofence_events_2026-03-01_2026-08-27.csv"
F_DEV = "telematics_device_assignments.csv"
F_UNITS = "fleet_units_2026-08.xlsx"
F_FMS = "fms_island_transactions_2025.csv"
F_WEST = "riverbend_direct_bill_2024-09_2026-08.csv"
F_CENTRAL = "crossroads_transactions_2024-09_2026-08.xlsx"
F_EAST = "smpm_card_activity_2024-09_2026-08.txt"
YARDS = ["Memphis", "Jackson", "Nashville", "Cookeville", "Knoxville", "Blountville"]
CODE = {"Memphis": "MEM", "Jackson": "JXN", "Nashville": "NSH", "Cookeville": "CKV", "Knoxville": "KNX",
        "Blountville": "TRI"}
ACCOUNT = {"Memphis": "West", "Jackson": "West", "Nashville": "Central", "Cookeville": "Central", "Knoxville": "East",
           "Blountville": "East"}
TZ = {"Memphis": CT, "Jackson": CT, "Nashville": CT, "Cookeville": CT, "Knoxville": ET, "Blountville": ET}
LOAD = 8000.0
BASE = (dt.date(2026, 6, 29), dt.date(2026, 8, 23))      # the eight complete weeks to Sunday 23 August
WIN = (dt.date(2026, 9, 14), dt.date(2026, 11, 8))
HOLIDAYS = {dt.date(2025, 5, 26), dt.date(2025, 7, 4), dt.date(2025, 9, 1), dt.date(2025, 11, 27),
            dt.date(2025, 11, 28), dt.date(2025, 12, 24), dt.date(2025, 12, 25), dt.date(2026, 1, 1),
            dt.date(2026, 5, 25), dt.date(2026, 7, 3), dt.date(2026, 9, 7), dt.date(2026, 11, 26),
            dt.date(2026, 11, 27)}


def working_days(a, b):
    n, d = 0, a
    while d <= b:
        n += d.weekday() < 5 and d not in HOLIDAYS
        d += dt.timedelta(days=1)
    return n


def floor_loads(g):
    return int(math.floor(g / LOAD + 1e-9))


# ------------------------------------------------------------------------------------------------- reading
class Pack:
    pass


def unit_from_name(name):
    for pat in (r"^D(\d{4}) ", r"^KSC-(\d{4})$", r"^HSS (\d{4}) "):
        m = re.match(pat, name)
        if m:
            return int(m.group(1))
    return None


def read(folder, cards=True, use_devices=True, dedupe_fms=True, day_tz=None):
    """use_devices, dedupe_fms and day_tz exist for the checks' mishandling tests (R2)."""
    P = Pack()
    P.folder = folder
    wb = openpyxl.load_workbook(os.path.join(folder, F_UNITS), read_only=True)
    ur = list(wb["Units"].iter_rows(values_only=True))
    sr = list(wb["Model Specs"].iter_rows(values_only=True))
    P.tank = {r[0]: float(r[2]) for r in sr[1:]}
    P.units = {}
    for r in ur[1:]:
        P.units[int(r[0])] = dict(yard=r[2], model=r[5], cls=r[6], status=r[8], tank=P.tank[r[5]])
    # device serials to vehicles (a new device reports under its serial until it is named)
    serial_unit = {}
    with open(os.path.join(folder, F_DEV), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            if use_devices:
                serial_unit[r["Device Serial"]] = unit_from_name(r["Vehicle"])

    def unit_of(name, serial):
        u = unit_from_name(name)
        return u if u is not None else serial_unit.get(serial)
    P.unit_of = unit_of
    # vehicle days
    P.days = collections.defaultdict(dict)
    with open(os.path.join(folder, F_DAYS), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            u = unit_of(r["Vehicle"], r["Device Serial"])
            if u is None:
                continue
            d = dt.date.fromisoformat(r["Day"])
            tz = day_tz or (CT if r["TZ"] == "CT" else ET)
            a = dt.datetime.combine(d, dt.time.fromisoformat(r["First Ignition"]), tz)
            b = dt.datetime.combine(d, dt.time.fromisoformat(r["Last Ignition Off"]), tz)
            P.days[u][d] = dict(a=a, b=b, fuel=float(r["Fuel Used (gal)"]), miles=float(r["Distance (mi)"]),
                                lat0=float(r["First Lat"]), lon0=float(r["First Lon"]), lat1=float(r["Last Lat"]),
                                lon1=float(r["Last Lon"]), odo0=float(r["Odometer Start"]),
                                odo1=float(r["Odometer End"]))
    # geofence crossings, paired into stays
    ev = collections.defaultdict(list)
    with open(os.path.join(folder, F_GEO), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            u = unit_of(r["Vehicle"], r["Device Serial"])
            if u is None:
                continue
            t = dt.datetime.strptime(r["Event Time"], "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=UTC)
            ev[u].append((t, r["Geofence"], r["Event"]))
    P.stays = {}
    P.unpaired = 0
    for u, L in ev.items():
        L.sort()
        open_ = {}
        S = []
        for (t, gname, kind) in L:
            if kind == "Enter":
                if gname in open_:
                    P.unpaired += 1
                open_[gname] = t
            else:
                t0 = open_.pop(gname, None)
                S.append((t0, t, gname))       # t0 None: in the geofence when the history starts
        for gname, t0 in open_.items():
            S.append((t0, None, gname))         # still inside at the end of the history
        S.sort(key=lambda x: (x[0] or x[1]))
        P.stays[u] = S
    # 2025 island transactions (the resent day's buffer appears twice under the same transaction numbers)
    P.fms = []
    seen = set()
    P.fms_dupes = 0
    with open(os.path.join(folder, F_FMS), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            if r["Txn"] + r["Site"] in seen:
                P.fms_dupes += 1
                if dedupe_fms:
                    continue
            seen.add(r["Txn"] + r["Site"])
            d = dt.datetime.strptime(r["Date"] + " " + r["Time"], "%m/%d/%Y %H:%M")
            P.fms.append(dict(unit=int(r["Unit"]), site=r["Site"][:3], t=d, gal=float(r["Gallons"])))
    if cards:
        P.cards = read_cards(folder)
    return P


def norm_unit(s):
    s = str(s).strip().upper().replace("HSS", "").replace(" ", "")
    return int(s) if s.isdigit() else None


def read_cards(folder):
    """Every card line as the three files give it: account, date and time as recorded, unit, product, gallons, and its
    role: a sale, a duplicate posting a later credit reverses ('dup'), or that credit ('credit')."""
    out = []
    with open(os.path.join(folder, F_WEST), encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    reversed_ = {r["REF_TRAN_ID"] for r in rows if r["REF_TRAN_ID"]}
    for r in rows:
        t = dt.datetime.strptime(r["TRAN_DATE_TIME"], "%Y-%m-%d %H:%M:%S")
        role = "credit" if r["REF_TRAN_ID"] else ("dup" if r["TRAN_ID"] in reversed_ else "sale")
        out.append(dict(acct="West", t=t, unit=norm_unit(r["UNIT_NO"]), diesel=r["PROD_CD"] == "01",
                        gal=float(r["QTY"]), site=r["SITE_NO"], city=r["SITE_CITY"].title(), st=r["SITE_ST"],
                        odo=int(r["ODOMETER"]), key=("W", r["TRAN_ID"]), role=role))
    wb = openpyxl.load_workbook(os.path.join(folder, F_CENTRAL), read_only=True)
    rows = list(wb["Transactions"].iter_rows(values_only=True))[1:]
    reversed_ = {r[20] for r in rows if r[20]}
    for r in rows:
        hh, mm = (int(x) for x in r[2].split(":"))
        t = dt.datetime.combine(r[1].date(), dt.time(hh, mm))
        role = "credit" if r[20] else ("dup" if r[0] in reversed_ else "sale")
        out.append(dict(acct="Central", t=t, unit=norm_unit(r[5]), diesel=r[12] == "Diesel", gal=float(r[13]),
                        site=r[8], city=r[10], st=r[11], odo=int(r[7]), key=("C", r[0]), role=role))
    seen = {}
    east = []
    with open(os.path.join(folder, F_EAST), encoding="utf-8") as fh:
        batch_week = None
        for ln in fh:
            ln = ln.rstrip("\r\n")
            if ln.startswith("H"):
                batch_week = ln[8:16]
            elif ln[:1] in ("D", "C"):
                body = ln[8:]                      # the record less its batch number
                sig = (batch_week, body)
                if sig in seen:                    # a week's batch transmitted again under the next batch number
                    continue
                seen[sig] = 1
                east.append(ln)
    credited = {ln[108:121] for ln in east if ln[0] == "C"}
    for ln in east:
        t = dt.datetime.strptime(ln[14:22] + ln[22:26], "%Y%m%d%H%M")
        sign = -1.0 if ln[0] == "C" else 1.0
        role = "credit" if ln[0] == "C" else ("dup" if ln[1:14] in credited else "sale")
        out.append(dict(acct="East", t=t, unit=norm_unit(ln[32:40]), diesel=ln[55:59] == "DSL2",
                        gal=sign * int(ln[59:68]) / 1000.0, site=ln[26:32].strip(), city=ln[82:106].strip().title(),
                        st=ln[106:108], odo=int(ln[48:55]), key=("E", ln[1:14]), role=role))
    return out


# ------------------------------------------------------------------------------------------------- the island draw
def reserve_share(P):
    """What a 2025 island put aboard a tractor back from a run that outlasted a tankful: its fills of three quarters of
    the tank or more, as a share of the tank."""
    fills = [f["gal"] / P.units[f["unit"]]["tank"] for f in P.fms
             if P.units[f["unit"]]["model"] == "T68SL" and f["gal"] >= 0.75 * P.units[f["unit"]]["tank"]]
    return 1.0 - sum(fills) / len(fills), len(fills)


def yard_stays(P, u):
    gname = "KSC %s Yard" % P.units[u]["yard"]
    return [(a, b) for (a, b, g) in P.stays.get(u, []) if g == gname]


def segments(P, u, start=dt.date(2026, 3, 1), split="time"):
    """The unit's driving time day by day from the start of the geofence history, cut at its yard stays:
    (start, end, gallons), in order."""
    tz = TZ[P.units[u]["yard"]]
    stays = [((a.astimezone(tz) if a else None), (b.astimezone(tz) if b else None)) for (a, b) in yard_stays(P, u)]
    lo = dt.datetime(1990, 1, 1, tzinfo=tz)
    hi = dt.datetime(2100, 1, 1, tzinfo=tz)
    bounds = [((s0 or lo), (s1 or hi)) for (s0, s1) in stays]
    out = []
    j = 0
    for d in sorted(P.days.get(u, {})):
        if d < start:
            continue
        r = P.days[u][d]
        a, b = r["a"], r["b"]
        pieces = [(a, b)]
        while j < len(bounds) and bounds[j][1] <= a:
            j += 1
        k = j
        while k < len(bounds) and bounds[k][0] < b:
            s0, s1 = bounds[k]
            k += 1
            nxt = []
            for (p0, p1) in pieces:
                if s1 <= p0 or s0 >= p1:
                    nxt.append((p0, p1))
                    continue
                if s0 > p0:
                    nxt.append((p0, s0))
                if s1 < p1:
                    nxt.append((s1, p1))
            pieces = nxt
        tot = sum((p1 - p0).total_seconds() for (p0, p1) in pieces)
        if tot <= 0:
            continue
        for (p0, p1) in pieces:
            share = (p1 - p0).total_seconds() / tot if split == "time" else 1.0 / len(pieces)
            out.append((p0, p1, r["fuel"] * share))
    return out, stays


def island_fills(P, u, cap, cut="visit", split="time", scale=1.0):
    """(arrival, gallons the island puts aboard) for each arrival at the yard: the stretch's fuel since the last
    departure from the yard, capped at the tankful. cut='overnight' cuts only at nights spent at the yard."""
    segs, stays = segments(P, u, split=split)
    arrivals = [s0 for (s0, s1) in stays if s0 is not None]
    if cut == "overnight":
        nights = set()
        for (s0, s1) in stays:
            if s0 is not None and (s1 is None or s1.date() > s0.date()):
                nights.add(s0)
        arrivals = [a for a in arrivals if a in nights]
    out = []
    i = 0
    acc = 0.0
    for a in sorted(arrivals):
        while i < len(segs) and segs[i][1] <= a:
            acc += segs[i][2] * scale
            i += 1
        out.append((a, min(acc, cap)))
        acc = 0.0
    return out


def draw_by_yard(P, base=BASE, cap_mode="reserve", cut="visit", r=None, scale=True, split="time", fuel_scale=None,
                 caps=None):
    """cap_mode: reserve (the tank less the 2025 reserve share), nominal (the tank), or caps (a unit -> gallons map)."""
    if r is None:
        r, _ = reserve_share(P)
    out = collections.defaultdict(float)
    for u, info in P.units.items():
        if info["status"] != "Active":
            continue
        cap = info["tank"] * (1.0 - r) if cap_mode == "reserve" else info["tank"]
        if caps is not None and u in caps:
            cap = caps[u]
        k_u = fuel_scale.get(u, 1.0) if fuel_scale else 1.0
        for (a, g) in island_fills(P, u, cap, cut, split, k_u):
            if base[0] <= a.date() <= base[1]:
                out[info["yard"]] += g
    k = working_days(*WIN) / working_days(*base) if scale else 1.0
    return {y: out[y] * k for y in YARDS}


def loads(draw):
    return {y: floor_loads(draw[y]) for y in YARDS}


# ------------------------------------------------------------------------------------------------- the rungs
def card_by_yard(P, a, b, diesel=True):
    out = collections.defaultdict(float)
    for f in P.cards:
        if diesel and not f["diesel"]:
            continue
        if f["role"] != "sale":
            continue
        if f["unit"] not in P.units or not (a <= f["t"].date() <= b):
            continue
        out[P.units[f["unit"]]["yard"]] += f["gal"]
    return {y: out[y] for y in YARDS}


def share_2025(P):
    """The 2025 islands' gallons against the same vehicles' engine fuel over the days each island dispensed, by yard,
    by model and pooled."""
    span = {}
    for f in P.fms:
        s = span.setdefault(f["site"], [f["t"].date(), f["t"].date()])
        s[0], s[1] = min(s[0], f["t"].date()), max(s[1], f["t"].date())
    isl = collections.defaultdict(float)
    for f in P.fms:
        isl[f["unit"]] += f["gal"]
    fuel = collections.defaultdict(float)
    for u, info in P.units.items():
        code = CODE[info["yard"]]
        if code not in span:
            continue
        a, b = span[code]
        fuel[u] = sum(r["fuel"] for d, r in P.days.get(u, {}).items() if a <= d <= b)
    by_yard, by_model = collections.defaultdict(lambda: [0.0, 0.0]), collections.defaultdict(lambda: [0.0, 0.0])
    for u in fuel:
        y = P.units[u]["yard"]
        m = P.units[u]["model"]
        by_yard[y][0] += isl[u]
        by_yard[y][1] += fuel[u]
        by_model[m][0] += isl[u]
        by_model[m][1] += fuel[u]
    pooled = sum(v[0] for v in by_yard.values()) / sum(v[1] for v in by_yard.values())
    return dict(span=span, yard={y: v[0] / v[1] for y, v in by_yard.items()},
                model={m: v[0] / v[1] for m, v in by_model.items()}, pooled=pooled, isl=dict(isl), fuel=dict(fuel))


def card_by_unit(P, a, b):
    out = collections.defaultdict(float)
    for f in P.cards:
        if f["diesel"] and f["role"] == "sale" and f["unit"] in P.units and a <= f["t"].date() <= b:
            out[f["unit"]] += f["gal"]
    return out


def near_yard(P, lat, lon, miles=1.0):
    import math as m
    best = None
    for y, (ya, yo) in YARD_LL.items():
        d = 3958.8 * 2 * m.asin(m.sqrt(m.sin(m.radians(lat - ya) / 2) ** 2 + m.cos(m.radians(lat)) *
                                       m.cos(m.radians(ya)) * m.sin(m.radians(lon - yo) / 2) ** 2))
        if d <= miles:
            best = y
    return best


YARD_LL = {"Memphis": (35.0412, -89.9631), "Jackson": (35.6352, -88.8713), "Nashville": (36.1107, -86.7202),
           "Cookeville": (36.1384, -85.4612), "Knoxville": (35.9783, -83.8842), "Blountville": (36.4687, -82.4021)}


def nights_at_yard_share(P, u, a, b):
    """Share of the unit's nights in [a, b] that it spent at its yard, read from the vehicle days' last positions."""
    y = P.units[u]["yard"]
    n = k = 0
    for d, r in P.days.get(u, {}).items():
        if a <= d <= b:
            n += 1
            k += near_yard(P, r["lat1"], r["lon1"]) == y
    return k / n if n else 0.0


def rungs(P, G=None):
    """Every route on file in the design note's ladder, each a dict yard -> gallons for the window."""
    R = {}
    wk = working_days(*WIN) / working_days(*BASE)
    c8 = card_by_yard(P, *BASE)
    R["R0 card gallons, trailing eight weeks"] = c8
    R["R0 card gallons x 40/39"] = {y: c8[y] * wk for y in YARDS}
    a13 = BASE[1] - dt.timedelta(days=7 * 13 - 1)
    c13 = card_by_yard(P, a13, BASE[1])
    R["R0 card gallons, 13 weeks"] = {y: c13[y] * working_days(*WIN) / working_days(a13, BASE[1]) for y in YARDS}
    ly = (WIN[0] - dt.timedelta(days=364), WIN[1] - dt.timedelta(days=364))
    R["R0 card gallons, same weeks of 2025"] = card_by_yard(P, *ly)
    S = share_2025(P)
    R["R1a pooled 2025 share"] = {y: c8[y] * S["pooled"] for y in YARDS}
    own = {"Memphis": "Jackson", "Jackson": "Jackson", "Knoxville": "Knoxville", "Blountville": "Knoxville"}
    R["R1a' 2025 share by account"] = {y: c8[y] * (S["yard"][own[y]] if y in own else S["pooled"]) for y in YARDS}
    cu = card_by_unit(P, *BASE)
    r1b = collections.defaultdict(float)
    for u, g in cu.items():
        m = P.units[u]["model"]
        r1b[P.units[u]["yard"]] += g * S["model"].get(m, 1.0)
    R["R1b 2025 share by model"] = {y: r1b[y] for y in YARDS}
    for k in ("R1a pooled 2025 share", "R1a' 2025 share by account", "R1b 2025 share by model"):
        R[k + " x 40/39"] = {y: R[k][y] * wk for y in YARDS}
    R["R2a tankful per run between nights at the yard"] = draw_by_yard(P, cut="overnight")
    r2b = collections.defaultdict(float)
    for u, g in cu.items():
        r2b[P.units[u]["yard"]] += g * nights_at_yard_share(P, u, *BASE)
    R["R2b card gallons x share of nights at the yard"] = {y: r2b[y] for y in YARDS}
    city = collections.defaultdict(float)
    for f in P.cards:
        if f["diesel"] and f["role"] == "sale" and f["unit"] in P.units and BASE[0] <= f["t"].date() <= BASE[1]:
            if f["city"] == P.units[f["unit"]]["yard"]:
                city[P.units[f["unit"]]["yard"]] += f["gal"]
    R["Reflex card fills in the yard's own city"] = {y: city[y] for y in YARDS}
    daily = collections.defaultdict(float)
    for u, g in cu.items():
        if nights_at_yard_share(P, u, *BASE) > 0.99:
            daily[P.units[u]["yard"]] += g
    R["Reflex daily-return vehicles only"] = {y: daily[y] for y in YARDS}
    gold = G or draw_by_yard(P)
    R["Reflex mean of the card gallons and the overnight cut"] = {
        y: (c8[y] + R["R2a tankful per run between nights at the yard"][y]) / 2 for y in YARDS}
    R["Golden"] = gold
    return R, S


def golden(P):
    r, n = reserve_share(P)
    draw = draw_by_yard(P, r=r)
    return dict(draw=draw, loads=loads(draw), total=sum(loads(draw).values()), reserve=r, n_reserve=n)
