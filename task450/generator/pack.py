"""Write task450's data files from the simulation: the two years' item-response exports, the item bank, the student
information system's enrollment extract and id merges, the platform's registration crosswalk, the campus
directory, two files the office keeps beside them, and the field dictionary."""
import csv
import datetime as dt

import numpy as np

import world as W

FILES = {
    "resp26": "grade5_math_responses_2026.csv",
    "resp25": "grade5_math_responses_2025.csv",
    "bank": "dma_item_bank_grade5_math.xlsx",
    "enroll": "sis_grade5_enrollment_2025-26.csv",
    "xwalk": "dma_registrations_2026.csv",
    "merges": "sis_student_id_merges_2026.csv",
    "directory": "campus_directory_2026-27.xlsx",
    "projection": "grade5_enrollment_projection_2026-27.csv",
    "learning": "math_professional_learning_hours_2025-26.csv",
    "dictionary": "data_dictionary.xlsx",
}
DISTRACTORS = ("projection", "learning")
FIX_POS = 4                      # the 2026 item whose key the bank carried wrongly
STRANDS = [("Number and operations", ["5.2A", "5.2B", "5.3A", "5.3E", "5.3G", "5.3K", "5.3L"]),
           ("Algebraic reasoning", ["5.4A", "5.4B", "5.4C", "5.4F", "5.4H"]),
           ("Geometry and measurement", ["5.5A", "5.6A", "5.6B", "5.7A", "5.8C"]),
           ("Data analysis", ["5.9A", "5.9C"]),
           ("Personal financial literacy", ["5.10A", "5.10E"])]


def _weekdays(a, b):
    d, out = a, []
    while d <= b:
        if d.weekday() < 5:
            out.append(d)
        d += dt.timedelta(days=1)
    return out


def bank_key(F):
    """The key as the item bank carries it: the 2026 item at FIX_POS was entered with a wrong letter."""
    key = list(F["key26"])
    true = key[FIX_POS - 1]
    key[FIX_POS - 1] = "C" if true != "C" else "D"
    return key


def items(F, rng):
    """One row per item in the bank: the 25 shared items, the 2025 form's own 25, the 2026 form's own 25."""
    rows = []
    tag = {}
    for p in range(1, W.N_ITEMS + 1):
        for iid in {F["ids25"][p - 1], F["ids26"][p - 1]}:
            if iid not in tag:
                s = STRANDS[int(rng.choice(5, p=[0.36, 0.26, 0.22, 0.10, 0.06]))]
                tag[iid] = (s[0], s[1][int(rng.integers(len(s[1])))])
    for p in F["drift"]:
        tag[F["ids25"][p - 1]] = W.DRIFT_TAG          # the three shared items on simplifying numerical expressions
    bk = bank_key(F)
    for p in range(1, W.N_ITEMS + 1):
        a, b = F["ids25"][p - 1], F["ids26"][p - 1]
        if a == b:
            rows.append(dict(item_id=a, strand=tag[a][0], teks=tag[a][1], key=F["key25"][p - 1], pos25=p, pos26=p,
                             first=2025))
        else:
            rows.append(dict(item_id=a, strand=tag[a][0], teks=tag[a][1], key=F["key25"][p - 1], pos25=p, pos26=None,
                             first=2025))
            rows.append(dict(item_id=b, strand=tag[b][0], teks=tag[b][1], key=bk[p - 1], pos25=None, pos26=p,
                             first=2026))
    rows.sort(key=lambda r: r["item_id"])
    return rows


def write_bank(F, path, rng):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Items"
    ws.append(["Item ID", "Strand", "TEKS", "Item type", "Key", "Points", "Position, 2025 form", "Position, 2026 form",
               "First used"])
    rows = items(F, rng)
    for r in rows:
        ws.append([r["item_id"], r["strand"], r["teks"], "Multiple choice", r["key"], 1, r["pos25"], r["pos26"],
                   r["first"]])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F4F6F")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, wd in zip("ABCDEFGHI", (11, 28, 8, 16, 6, 7, 11, 11, 9)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:I%d" % (len(rows) + 1)
    wb.save(path)
    return len(rows)


def _abbrev(name):
    n = name.upper().replace(" ELEMENTARY", " EL")
    for a, b in (("AVENUE", "AVE"), ("STREET", "ST"), ("ROAD", "RD"), ("TERRACE", "TER"), ("CROSSING", "XING"),
                 ("PRAIRIE", "PR"), ("HOLLOW", "HOLW"), ("MEADOW", "MDW"), ("JUNIPER", "JUNP")):
        n = n.replace(a, b)
    return n


def write_resp25(w, path, rng):
    """The old reporting module's export: one row per 2025 test, answers in position columns, the platform's campus
    code and abbreviated campus name as recorded in spring 2025."""
    A, C = w["A"], w["C"]
    days = _weekdays(*W.WINDOW_2025)
    n = len(A["campus"])
    ids = rng.choice(np.arange(1000000, 1999999), size=n, replace=False)
    tests = rng.choice(np.arange(100000, 999999), size=n, replace=False)
    order = np.argsort(ids)
    with open(path, "w", newline="") as fh:
        wr = csv.writer(fh)
        wr.writerow(["TEST_ID", "STUDENT_LOCAL_ID", "SCHOOL_CODE", "SCHOOL_NAME", "GRADE", "FORM", "TEST_DATE"] +
                    ["I%02d" % p for p in range(1, W.N_ITEMS + 1)])
        for i in order:
            c = A["campus"][i]
            d = days[int(rng.integers(len(days)))]
            wr.writerow(["D25-%06d" % tests[i], "%07d" % ids[i], C["legacy"][c], _abbrev(c), "05", W.FORM25,
                         d.strftime("%m/%d/%Y")] + list(A["resp"][i]))
    return n


def tested_rows(w):
    B = w["B"]
    return [i for i, st in enumerate(B["stu"]) if st["tested"]]


def write_resp26(w, path, rng):
    """The new reporting module's export: one row per answer, by item id, with the session and the submit time."""
    B, F = w["B"], w["F"]
    days = _weekdays(*W.WINDOW)
    rows = 0
    idx = sorted(tested_rows(w), key=lambda i: B["stu"][i]["vendor"])
    with open(path, "w", newline="") as fh:
        wr = csv.writer(fh)
        wr.writerow(["test_taker_id", "form_code", "session", "item_position", "item_id", "response", "submitted_at"])
        for i in idx:
            st = B["stu"][i]
            d1 = days[int(rng.integers(len(days) - 1))]
            d2 = days[min(len(days) - 1, days.index(d1) + int(rng.integers(1, 3)))]
            for sess, d in ((1, d1), (2, d2)):
                t0 = dt.datetime.combine(d, dt.time(8, 30)) + dt.timedelta(minutes=int(rng.integers(0, 150)))
                for p in range(1 + 25 * (sess - 1), 26 + 25 * (sess - 1)):
                    t = t0 + dt.timedelta(seconds=int(40 + 75 * (p - 25 * (sess - 1)) + rng.integers(0, 50)))
                    wr.writerow([st["vendor"], W.FORM26, sess, p, F["ids26"][p - 1], B["resp"][i][p - 1],
                                 t.strftime("%Y-%m-%d %H:%M:%S")])
                    rows += 1
    return rows


EXIT = {"moved": "TW", "winter": "TW", "left": "LD"}


def spells_table(w):
    """Every enrollment spell in the new system's 2025-26 history: (student_id, campus, entry, exit, exit_code)."""
    B, C = w["B"], w["C"]
    out = []
    for st in B["stu"]:
        for c, a, b in st["spells"]:
            code = EXIT.get(st["kind"], "EY") if b is not None else "EY"
            out.append((st["id"], c, a, b or W.YEAR_END, code))
        if "ghost" in st:
            c, a, _ = st["ghost_spell"]
            out.append((st["ghost"], c, a, W.YEAR_END, "EY"))
    out.sort(key=lambda r: (r[0], r[2]))
    return out


def write_enroll(w, path):
    C = w["C"]
    rows = spells_table(w)
    with open(path, "w", newline="") as fh:
        wr = csv.writer(fh)
        wr.writerow(["student_id", "campus_number", "grade", "entry_date", "exit_date", "exit_code"])
        for sid, c, a, b, code in rows:
            wr.writerow([sid, C["new"][c], "05", a.isoformat(), b.isoformat(), code])
    return len(rows)


def write_xwalk(w, path, rng):
    """The platform's registrations: the roster load of 23 March 2026 (each student at the campus enrolled on that
    day) and the students who enrolled after it (on their entry date)."""
    B, C = w["B"], w["C"]
    load = dt.date(2026, 3, 23)
    rows = []
    for st in B["stu"]:
        on = [c for c, a, b in st["spells"] if a <= load and (b is None or b >= load)]
        if on:
            rows.append((st["vendor"], st["id"], C["new"][on[0]], load))
        else:
            later = sorted((a, c) for c, a, b in st["spells"] if load < a <= W.WINDOW[1])
            if later:
                rows.append((st["vendor"], st["id"], C["new"][later[0][1]], later[0][0]))
        if "ghost" in st:
            rows.append((st["ghost_vendor"], st["ghost"], C["new"][st["ghost_spell"][0]], load))
    rows.sort()
    with open(path, "w", newline="") as fh:
        wr = csv.writer(fh)
        wr.writerow(["test_taker_id", "student_id", "campus_number", "registered_on"])
        for v, sid, c, d in rows:
            wr.writerow([v, sid, c, d.isoformat()])
    return len(rows)


def write_merges(w, path):
    B = w["B"]
    rows = sorted(B["merges"], key=lambda r: (r[2], r[0]))
    with open(path, "w", newline="") as fh:
        wr = csv.writer(fh)
        wr.writerow(["merged_id", "surviving_id", "merged_on", "reason"])
        for a, b, d in rows:
            wr.writerow([a, b, d.isoformat(), "Duplicate person record"])
    return len(rows)


def write_directory(w, path, rng):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    C = w["C"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Campuses"
    ws.append(["Campus number", "Campus", "Learning community", "Grades", "Opened", "Status, 2026-27"])
    for c in sorted(C["names"], key=lambda x: C["new"][x]):
        opened = 2025 if c == "Vereda Elementary" else int(1951 + round(int(C["legacy"][c]) - 101) * 1.25)
        status = "Open"
        if c == W.CLOSED[0]:
            status = "Closed 30 June 2026; attendance zone to %s" % W.CLOSED[1]
        grades = "K-5" if (C["names"].index(c) * 31) % 5 else "PK-5"
        ws.append([C["new"][c], c, C["community"][c], grades, opened, status])
    for col in ws[1]:
        col.font = Font(bold=True, color="FFFFFF")
        col.fill = PatternFill("solid", fgColor="2F4F6F")
        col.alignment = Alignment(wrap_text=True, vertical="top")
    for col, wd in zip("ABCDEF", (10, 32, 12, 7, 8, 52)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A2"
    wb.save(path)
    return len(C["names"])


def write_projection(w, path, rng):
    """The demographer's projected grade-5 enrollment for 2026-27, by campus (planning, not results)."""
    C = w["C"]
    with open(path, "w", newline="") as fh:
        wr = csv.writer(fh)
        wr.writerow(["campus_number", "campus_name", "projected_grade5_2026_27", "method"])
        for c in sorted(C["names"], key=lambda x: C["new"][x]):
            if c == W.CLOSED[0]:
                continue
            base = C["size"][c] + (C["size"][W.CLOSED[0]] if c == W.CLOSED[1] else 0)
            wr.writerow([C["new"][c], c, int(round(base * rng.uniform(0.93, 1.06))), "cohort survival, 3-year"])
    return len(C["names"]) - 1


def write_learning(w, path, rng):
    """Hours of mathematics professional learning grade-3 to grade-5 teachers logged in 2025-26, by campus and
    month (the curriculum office's log)."""
    C = w["C"]
    months = [dt.date(2025, m, 1) for m in range(8, 13)] + [dt.date(2026, m, 1) for m in range(1, 6)]
    n = 0
    with open(path, "w", newline="") as fh:
        wr = csv.writer(fh)
        wr.writerow(["campus_number", "month", "teachers", "hours", "sessions"])
        for c in sorted(C["names"], key=lambda x: C["new"][x]):
            t = max(3, int(round(C["size"][c] / 22.0 * 3)))
            for m in months:
                s = int(rng.integers(0, 4))
                h = round(s * t * rng.uniform(0.8, 1.6), 1)
                wr.writerow([C["new"][c], m.strftime("%Y-%m"), t, h, s])
                n += 1
    return n


FIELDS = [
    (FILES["resp26"], "test_taker_id", "The platform's identifier for a registered test taker"),
    (FILES["resp26"], "form_code", "Test form"),
    (FILES["resp26"], "session", "Testing session (positions 1 to 25 in session 1, 26 to 50 in session 2)"),
    (FILES["resp26"], "item_position", "Position of the item on the form"),
    (FILES["resp26"], "item_id", "Item bank identifier"),
    (FILES["resp26"], "response", "Answer chosen: A to D; blank when no answer was given; * when more than one was marked"),
    (FILES["resp26"], "submitted_at", "Time the answer was submitted, district local time"),
    (FILES["resp25"], "TEST_ID", "The platform's identifier for the test"),
    (FILES["resp25"], "STUDENT_LOCAL_ID", "Student identifier in the district's records at the time of testing"),
    (FILES["resp25"], "SCHOOL_CODE", "Campus code as the platform recorded it"),
    (FILES["resp25"], "SCHOOL_NAME", "Campus name as the platform recorded it"),
    (FILES["resp25"], "GRADE", "Grade tested"),
    (FILES["resp25"], "FORM", "Test form"),
    (FILES["resp25"], "TEST_DATE", "Date tested (MM/DD/YYYY)"),
    (FILES["resp25"], "I01 to I50", "Answer chosen at each position of the form: A to D; blank when no answer was given; * when more than one was marked"),
    (FILES["bank"], "Item ID", "Item bank identifier"),
    (FILES["bank"], "Key", "The correct answer"),
    (FILES["bank"], "Position, 2025 form / 2026 form", "Where the item appears on each grade-5 form; blank when it is not on that form"),
    (FILES["bank"], "First used", "The year the item was first administered"),
    (FILES["enroll"], "student_id", "Student identifier in the student information system"),
    (FILES["enroll"], "campus_number", "Campus number (see the campus directory)"),
    (FILES["enroll"], "entry_date / exit_date", "First and last day of the enrollment at the campus (the last day of school for enrollments open at year end)"),
    (FILES["enroll"], "exit_code", "EY end of year; TW transferred to another district campus; LD left the district"),
    (FILES["xwalk"], "test_taker_id", "The platform's identifier for a registered test taker"),
    (FILES["xwalk"], "student_id", "Student identifier in the student information system"),
    (FILES["xwalk"], "campus_number", "Campus the student was registered at"),
    (FILES["xwalk"], "registered_on", "Date of registration (the roster load of 23 March 2026, or enrollment after it)"),
    (FILES["merges"], "merged_id / surviving_id", "A student id retired into the id kept for the same student"),
    (FILES["merges"], "merged_on", "Date of the merge"),
    (FILES["directory"], "Campus number", "The district's campus number"),
    (FILES["directory"], "Learning community", "The learning community the campus belongs to"),
    (FILES["directory"], "Status, 2026-27", "Open, or the closure and where the attendance zone went"),
    (FILES["projection"], "projected_grade5_2026_27", "The demographer's projection of grade-5 enrollment for 2026-27"),
    (FILES["learning"], "hours", "Professional learning hours logged by the campus's grade-3 to grade-5 mathematics teachers"),
]


def write_dictionary(path):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fields"
    ws.append(["File", "Field", "Meaning"])
    for f in FIELDS:
        ws.append(list(f))
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F4F6F")
    for col, wd in zip("ABC", (40, 32, 96)):
        ws.column_dimensions[col].width = wd
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    wb.save(path)
    return len(FIELDS)
