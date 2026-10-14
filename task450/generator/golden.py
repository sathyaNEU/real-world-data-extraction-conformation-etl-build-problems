"""The golden deliverables, computed from the shipped files only: intervention_schools_2027.docx (the office's memo to
the superintendent's cabinet, with the campus chart) and grade5_math_results_2026.xlsx (every rated campus, the 2025
test by learning community, the students with no 2026 result); and figures.json, every figure the write-up states."""
import io
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import numpy as np      # noqa: E402

import calc as C        # noqa: E402
import reader as RD     # noqa: E402
import world as W       # noqa: E402

DOCX = "intervention_schools_2027.docx"
XLSX = "grade5_math_results_2026.xlsx"
GOLD_METHOD = "chained equipercentile"


def fmt(n):
    return "{:,}".format(int(n))


def pct(k, n):
    return "%.1f" % (100.0 * k / n)


def ask_2025(L):
    """Each learning community's 2025 test takers and those meeting the standard (raw 27 or more on the 2025 form),
    each test placed by the campus name the old platform recorded."""
    D = L["directory"]
    by_name = {v["name"].upper(): k for k, v in D.items()}
    raw = L["S25"].sum(1)
    out = {}
    for nm, r in zip(L["name25"], raw):
        c = by_name[RD.expand(nm)]
        k = D[L["into"].get(c, c)]["community"]
        n, m = out.get(k, (0, 0))
        out[k] = (n + 1, m + int(r >= 27))
    return out


def ask_2025_by_code(L):
    """The same, each test placed by its campus code read as a 2026 campus number."""
    D = L["directory"]
    raw = L["S25"].sum(1)
    out = {}
    for code, r in zip(L["code25"], raw):
        k = D[L["into"].get(code, code)]["community"]
        n, m = out.get(k, (0, 0))
        out[k] = (n + 1, m + int(r >= 27))
    return out


def ask_untested(L, merged=True):
    """Grade-5 students enrolled at some time in the 2026 testing window with no 2026 result, by learning community:
    one person per id after the merges, a person with any test result left out."""
    D = L["directory"]
    person = (lambda s: L["merges"].get(s, s)) if merged else (lambda s: s)
    tested = {person(L["xw"][t]) for t in L["takers"]}
    places = {}
    for sid, sp in L["spells"].items():
        for c, a, b, code in sp:
            if a <= W.WINDOW[1] and b >= W.WINDOW[0]:
                places.setdefault(person(sid), set()).add(D[L["into"].get(c, c)]["community"])
    out = {}
    for p, ks in places.items():
        if p in tested:
            continue
        assert len(ks) == 1, (p, ks)
        k = ks.pop()
        out[k] = out.get(k, 0) + 1
    return out


def compute(target):
    L = RD.load(target)
    scr = C.screened(L["S25"], L["S26"], L["anchor"])
    out, rg, info = C.methods(L["S25"], L["S26"], scr, rasch=False)
    r = C.passing(out[GOLD_METHOD])
    blind, _, _ = C.methods(L["S25"], L["S26"], L["anchor"], rasch=True)
    rb = sorted({C.passing(f) for f in blind.values()})
    tab = RD.campus_table(L, r)
    des = RD.designated(tab)
    X = dict(L=L, r=r, tab=tab, des=des, info=info,
             r_identity=C.passing(rg["identity (the 2025 conversion)"]),
             r_rg=C.passing(rg["random groups equipercentile"]),
             a25=ask_2025(L), un=ask_untested(L))
    X["des_identity"] = RD.designated(RD.campus_table(L, X["r_identity"]))
    X["des_rg"] = RD.designated(RD.campus_table(L, X["r_rg"]))
    X["scr"] = scr
    X["flagged"] = [j for j in L["anchor"] if j not in scr]
    X["r_blind"] = rb
    X["des_blind"] = {x: RD.designated(RD.campus_table(L, x)) for x in rb}
    X["V1"] = float(L["S25"][:, scr].sum(1).mean())
    X["V2"] = float(L["S26"][:, scr].sum(1).mean())
    X["p_flag"] = [(float(L["S25"][:, j].mean()), float(L["S26"][:, j].mean())) for j in X["flagged"]]
    X["p_kept"] = (float(L["S25"][:, scr].mean()), float(L["S26"][:, scr].mean()))
    X["X"] = float(L["S25"].sum(1).mean())
    X["Y"] = float(L["S26"].sum(1).mean())
    return X


def name(L, c):
    return L["directory"][c]["name"]


def pct1(p):
    return "%.1f" % (100.0 * p)


def figures(X):
    L, tab = X["L"], X["tab"]
    des = sorted(X["des"], key=lambda c: tab[c][1] / tab[c][0])
    tn = sum(n for n, k in tab.values())
    tk = sum(k for n, k in tab.values())
    F = dict(
        passing_raw=str(X["r"]),
        passing_raw_identity=str(X["r_identity"]),
        passing_raw_rg=str(X["r_rg"]),
        designated_count=str(len(des)),
        designated_identity=str(len(X["des_identity"])),
        designated_rg=str(len(X["des_rg"])),
        designated=[name(L, c) for c in des],
        designated_rates={name(L, c): pct(tab[c][1], tab[c][0]) for c in des},
        anchor_2025="%.2f" % X["V1"], anchor_2026="%.2f" % X["V2"], anchor_gain="%.2f" % (X["V2"] - X["V1"]),
        raw_2025="%.2f" % X["X"], raw_2026="%.2f" % X["Y"],
        shared_items=str(len(L["anchor"])),
        kept_items=str(len(X["scr"])),
        flagged_positions=[str(j + 1) for j in X["flagged"]],
        flagged_correct_2025=[pct1(a) for a, b in X["p_flag"]],
        flagged_correct_2026=[pct1(b) for a, b in X["p_flag"]],
        kept_correct_2025=pct1(X["p_kept"][0]), kept_correct_2026=pct1(X["p_kept"][1]),
        flagged_teks=W.DRIFT_TAG[1],
        passing_raw_blind=[str(x) for x in X["r_blind"]],
        designated_blind=[str(len(X["des_blind"][x])) for x in X["r_blind"]],
        district_rate=pct(tk, tn), district_tested=fmt(tn), district_meeting=fmt(tk),
        rated_campuses=str(len(tab)),
        campuses={name(L, c): dict(tested=fmt(n), meeting=fmt(k), rate=pct(k, n))
                  for c, (n, k) in sorted(tab.items(), key=lambda kv: name(L, kv[0]))},
        ask_2025={k: dict(tested=fmt(X["a25"][k][0]), rate=pct(X["a25"][k][1], X["a25"][k][0])) for k in W.COMMUNITIES},
        ask_untested={k: fmt(X["un"].get(k, 0)) for k in W.COMMUNITIES},
        ask_untested_district=fmt(sum(X["un"].values())),
        key_item=L["notice"]["item"], key_wrong=L["notice"]["wrong"], key_right=L["notice"]["right"],
        tests_2025=fmt(len(L["S25"])), tests_2026=fmt(len(L["S26"])), passing_raw_below=str(X["r"] - 1),
        unrated={name(L, c): str(n) for c, (n, k) in RD.campus_table(L, X["r"], n_min=1).items() if n < W.N_MIN},
        scale_intercept="141", scale_slope="6", scale_standard_raw="26.5", report_sections=["3", "4.1"],
        clauses=["2", "3", "4", "5", "6", "7"],
        line="40", standard="300", n_min="10", scale_formula="141 + 6",
    )
    return F


def chart_png(X):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.ticker import FuncFormatter
    L, tab = X["L"], X["tab"]
    order = sorted(tab, key=lambda c: tab[c][1] / tab[c][0])
    vals = [100.0 * tab[c][1] / tab[c][0] for c in order]
    cols = ["#7a2333" if c in X["des"] else "#a9b4bd" for c in order]
    fig, ax = plt.subplots(figsize=(6.6, 3.3), dpi=150)
    ax.bar(range(len(order)), vals, color=cols, width=0.78)
    ax.axhline(40, color="#1f1f1f", linewidth=0.9, linestyle=(0, (4, 2)))
    ax.text(len(order) - 0.5, 41.2, "Designation line, 40%", ha="right", va="bottom", fontsize=7.4, color="#1f1f1f")
    ax.set_xticks([])
    ax.set_xlim(-0.8, len(order) - 0.2)
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: "%d%%" % v))
    ax.set_ylabel("Meeting the standard (full academic year)", fontsize=7.8)
    ax.set_xlabel("The %d rated elementary campuses, lowest to highest" % len(order), fontsize=7.8)
    ax.set_title("%d campuses fall under the 40 per cent line (dark bars)" % len(X["des"]), fontsize=9.2, loc="left")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(labelsize=7.5)
    ax.set_ylim(0, max(vals) * 1.1)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def _page_field(paragraph):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def write_docx(X, F, path):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    L, tab = X["L"], X["tab"]
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.9)
        sec.top_margin = sec.bottom_margin = Inches(0.8)
    h = doc.add_paragraph()
    rh = h.add_run("%s  |  %s" % (W.DISTRICT, W.OFFICE))
    rh.font.size = Pt(9)
    rh.font.color.rgb = RGBColor(0x5C, 0x5C, 0x5C)
    foot = doc.sections[0].footer.paragraphs[0]
    fr = foot.add_run("%s  |  2027 mathematics intervention coaching: designated campuses  |  Page " % W.SHORT)
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(0x5C, 0x5C, 0x5C)
    _page_field(foot)
    t = doc.add_paragraph()
    rt = t.add_run("%s campuses fall under the 40 per cent line on the spring 2026 grade-5 results" % (
        {15: "Fifteen"}.get(len(X["des"]), str(len(X["des"])))))
    rt.bold = True
    rt.font.size = Pt(14)
    rt.font.color.rgb = RGBColor(0x7A, 0x23, 0x33)
    meta = doc.add_paragraph()
    meta.add_run("To: %s, %s, and the superintendent's cabinet\nFrom: %s, %s\nDate: 12 October 2026\n"
                 "Subject: campuses designated for 2027 mathematics intervention coaching (District Improvement Plan, "
                 "2027 amendment)" % (W.PEOPLE["superintendent"][0], W.PEOPLE["superintendent"][1],
                                      W.PEOPLE["director"][0], W.PEOPLE["director"][1])).font.size = Pt(9.5)
    p = doc.add_paragraph()
    r1 = p.add_run("Place the 2027 mathematics intervention coaches at the %s campuses below." % F["designated_count"])
    r1.bold = True
    p.add_run(" Each is a rated elementary campus where fewer than 40 per cent of its tested full-academic-year grade-5 "
              "students met the standard on the spring 2026 District Mathematics Assessment (designation rules 2 to "
              "7). Across the %s rated campuses, %s per cent met it." % (F["rated_campuses"], F["district_rate"]))
    tbl = doc.add_table(rows=1, cols=5)
    tbl.style = "Light List Accent 1"
    for i, hd in enumerate(("Campus", "Learning community", "Tested", "Met the standard", "Share (%)")):
        tbl.rows[0].cells[i].text = hd
    for nm in F["designated"]:
        c = [k for k, v in L["directory"].items() if v["name"] == nm][0]
        n, k = tab[c]
        row = tbl.add_row().cells
        row[0].text = nm
        row[1].text = L["directory"][c]["community"]
        row[2].text = fmt(n)
        row[3].text = fmt(k)
        row[4].text = F["designated_rates"][nm]
    src = doc.add_paragraph()
    rs = src.add_run("Source: spring 2026 responses (platform export of 9 October 2026) scored with the corrected key; "
                     "student information system enrollment for 2025-26. Tested and met are full-academic-year students.")
    rs.italic = True
    rs.font.size = Pt(8)
    doc.add_picture(io.BytesIO(chart_png(X)), width=Inches(6.3))
    doc.add_paragraph(
        "Placing 2026 scores on the scale. The 2026 grade-5 form keeps %s of the 2025 form's items in the same "
        "positions and replaces the other %s with new ones, so the 2025 conversion (141 + 6 x raw score) belongs to "
        "the 2025 form only. On the %s shared items the 2026 students averaged %s points against the 2025 students' "
        "%s: this cohort is stronger, while its average raw score on the whole form fell from %s to %s, so the new "
        "items are harder by more than the cohort gained. Linked to the 2025 scale through those shared items, a 2026 "
        "raw score of %s converts to 300 or more and %s to less, so %s is the passing raw score on the 2026 form."
        % (F["shared_items"], 50 - int(F["shared_items"]), F["kept_items"], F["anchor_2026"], F["anchor_2025"],
           F["raw_2025"], F["raw_2026"], F["passing_raw"], int(F["passing_raw"]) - 1, F["passing_raw"]))
    doc.add_paragraph(
        "Which shared items carry the link. Before linking, each shared item's difficulty was compared across the two "
        "years. Items %s, %s and %s, all written to TEKS %s (simplifying numerical expressions), moved far more than "
        "the rest: the 2026 students answered them correctly %s, %s and %s per cent of the time against %s, %s and %s "
        "per cent in 2025, while the other %s rose from %s to %s per cent on average. The new curriculum reaches that "
        "standard a year earlier, so those three items no longer measure the same thing in both years, and the link "
        "runs through the other %s. Every standard screen flags the same three (the change in Rasch difficulty, "
        "Mantel-Haenszel, the delta plot)."
        % (tuple(F["flagged_positions"]) + (F["flagged_teks"],) + tuple(F["flagged_correct_2026"]) +
           tuple(F["flagged_correct_2025"]) + (F["kept_items"], F["kept_correct_2025"], F["kept_correct_2026"],
                                               F["kept_items"])))
    doc.add_paragraph(
        "Why not the published conversion. Read through the 2025 conversion, the passing raw score stays at %s and %s "
        "campuses fall under the line; matching the 2026 score distribution to 2025's, which assumes the two cohorts "
        "are alike, puts it at %s and designates %s. The shared items show the cohorts are not alike: the 2025 conversion "
        "counts the new items' difficulty against every student, and matching the distributions reads this cohort's gain "
        "as an easier form. Linked through all %s shared items, the three items' rise is credited to the whole cohort: "
        "the passing raw score falls to %s or below, depending on the method, and %s or fewer campuses are designated."
        % (F["passing_raw_identity"], F["designated_identity"], F["passing_raw_rg"], F["designated_rg"],
           F["shared_items"], F["passing_raw_blind"][-1], F["designated_blind"][-1]))
    notes = doc.add_paragraph()
    rn = notes.add_run(
        "Notes. Item %s was scored with key %s, as the notice of 19 May 2026 corrects the bank's %s. Delgadillo "
        "Elementary's students are counted with Tobar Elementary, which takes in its zone for 2026-27 (rule 5). Vereda "
        "Elementary, opened in August 2025, had fewer than ten full-academic-year students tested and is not rated "
        "(rule 6). Shares are rounded to one decimal; designation compares the unrounded share with 40 per cent. "
        "Every rated campus's figures are in grade5_math_results_2026.xlsx."
        % (F["key_item"], F["key_right"], F["key_wrong"]))
    rn.font.size = Pt(8.5)
    rn.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    doc.save(path)


def write_xlsx(X, F, path):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    L, tab = X["L"], X["tab"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Campus results"
    ws.append(["Spring 2026 grade-5 mathematics, rated elementary campuses (full-academic-year students)"])
    ws.append([])
    ws.append(["Campus number", "Campus", "Learning community", "Tested", "Met the standard", "Share (%)",
               "Designated for 2027"])
    for c in sorted(tab, key=lambda x: name(L, x)):
        n, k = tab[c]
        ws.append([int(c), name(L, c), L["directory"][c]["community"], n, k, round(100.0 * k / n, 1),
                   "Yes" if c in X["des"] else None])
    last = ws.max_row
    tn = sum(n for n, k in tab.values())
    tk = sum(k for n, k in tab.values())
    ws.append([None, "All rated campuses", None, tn, tk, round(100.0 * tk / tn, 1), len(X["des"])])
    for col, wd in zip("ABCDEFG", (10, 30, 12, 9, 11, 10, 12)):
        ws.column_dimensions[col].width = wd
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="7A2333")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(bold=True, size=12)
    for row in ws.iter_rows(min_row=4, max_row=last + 1):
        row[3].number_format = "#,##0"
        row[4].number_format = "#,##0"
        row[5].number_format = "0.0"
    for c in ws[last + 1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "C4"
    ws.print_title_rows = "3:3"
    ws2 = wb.create_sheet("2025 by community")
    ws2.append(["Spring 2025 grade-5 mathematics test takers by learning community (2025 form)"])
    ws2.append([])
    ws2.append(["Learning community", "Took the 2025 test", "Met the standard (%)"])
    for k in W.COMMUNITIES:
        n, m = X["a25"][k]
        ws2.append([k, n, round(100.0 * m / n, 1)])
    for col, wd in zip("ABC", (20, 18, 20)):
        ws2.column_dimensions[col].width = wd
    for c in ws2[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="7A2333")
    ws2["A1"].font = Font(bold=True, size=12)
    for row in ws2.iter_rows(min_row=4):
        row[1].number_format = "#,##0"
        row[2].number_format = "0.0"
    ws3 = wb.create_sheet("No 2026 result")
    ws3.append(["Grade-5 students enrolled at some time in the testing window (20 April to 1 May 2026) with no 2026 "
                "result"])
    ws3.append([])
    ws3.append(["Learning community", "Students"])
    for k in W.COMMUNITIES:
        ws3.append([k, X["un"].get(k, 0)])
    ws3.append(["District", sum(X["un"].values())])
    ws3.column_dimensions["A"].width = 20
    ws3.column_dimensions["B"].width = 12
    for c in ws3[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="7A2333")
    ws3["A1"].font = Font(bold=True, size=12)
    for c in ws3[ws3.max_row]:
        c.font = Font(bold=True)
    ws4 = wb.create_sheet("Notes")
    notes = [
        "Shares are full-academic-year students who tested and met the standard over full-academic-year students who "
        "tested (designation rules 3 and 4). A student meets the standard at a scale score of 300 or above on the "
        "reporting scale set on the 2025 grade-5 form (rule 2).",
        "The 2026 form shares %s items with the 2025 form; 2026 scores are placed on the scale through those items, "
        "and a raw score of %s or more on the 2026 form meets the standard." % (F["shared_items"], F["passing_raw"]),
        "Campuses are as organised for 2026-27: Delgadillo Elementary's students count with Tobar Elementary (rule 5). "
        "Vereda Elementary had fewer than ten full-academic-year students tested and is not rated (rule 6).",
        "2025 by community: every 2025 grade-5 test taker, scored with the 2025 key; a raw score of 27 or more on the "
        "2025 form meets the standard. Each test is placed by the campus the 2025 platform recorded, under its name: "
        "the student information system renumbered every campus in January 2026.",
        "No 2026 result: students enrolled at a district elementary campus in grade 5 on any day from 20 April to "
        "1 May 2026 with no 2026 test, one person per student after the system's id merges of August 2026.",
        "Prepared by the Office of Assessment and Accountability, 12 October 2026, from the extract of 9 October 2026.",
    ]
    ws4.append(["Notes"])
    for n_ in notes:
        ws4.append([n_])
    ws4.column_dimensions["A"].width = 120
    for row in ws4.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    ws4["A1"].font = Font(bold=True, size=12)
    wb.save(path)


def main(target, golden, figures_path):
    os.makedirs(golden, exist_ok=True)
    X = compute(target)
    F = figures(X)
    write_docx(X, F, os.path.join(golden, DOCX))
    write_xlsx(X, F, os.path.join(golden, XLSX))
    import postprocess as PP
    by = W.PEOPLE["director"][0]
    PP.normalize_ooxml(os.path.join(golden, DOCX), by, by, "2026-10-12T14:05:00Z", "2026-10-12T19:26:44Z")
    PP.normalize_ooxml(os.path.join(golden, XLSX), by, by, "2026-10-12T13:40:00Z", "2026-10-12T19:12:08Z", excel=True)
    stamp = 1791837000.0
    for f in os.listdir(golden):
        os.utime(os.path.join(golden, f), (stamp, stamp))
    with open(figures_path, "w") as fh:
        json.dump(F, fh, indent=1, ensure_ascii=False)
        fh.write("\n")
    return X
