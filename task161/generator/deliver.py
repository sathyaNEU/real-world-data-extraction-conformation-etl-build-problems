"""task161 goldens: the call and the asks computed from target/ alone (read.py, calc.py, asks.py), written
as the planner sends them (overflow_blocks_h1_2027.xlsx, port_share_2026.pdf), and generator/figures.json
for figures_check.py and the write-up."""
from __future__ import annotations

import datetime as dt
import json
import os

import numpy as np

import asks
import calc
import postprocess as PP
import read

SEASONAL = ["PD2047", "PD2052", "PD2058", "PD2102", "PD2108"]
MONTHS = ["", "January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
          "November", "December"]


def compute(target):
    B = read.book(target)
    g, parts = calc.golden(B, return_parts=True)
    inp = read.ask_inputs(target)
    teu, share = asks.a2_golden(inp["receipts"], inp["manifest"], inp["LA"], inp["LB"], inp["to_local"])
    fade = calc.rival_decay(B)
    rivals = {"fade": fade, "book_burn": calc.rival_book_burn(B, "slope"),
              "book_burn_totals": calc.rival_book_burn(B, "total"),
              "per_account_fall": calc.rival_per_account_monthly_fall(B), "no_payback": calc.rival_no_payback(B)}
    return B, g, parts, inp, teu, share, rivals


def flat_hold(B):
    """Today's stock held flat through the term: the close of 6 December, November's mean, December to date."""
    tot = sum(B.acc[c]["oh"] for c in B.current)
    nov = [B.idx[d] for d in calc.month_days((2026, 11))]
    dec = [B.idx[d] for d in calc.month_days((2026, 12)) if d in B.idx]
    out = {}
    for name, v in (("close", tot[B.idx[B.last]]), ("november", tot[nov].mean()), ("december", tot[dec].mean())):
        out[name] = calc.blocks([float(v)] * 6)["blocks"]
    return out


def figures(B, g, parts, inp, teu, share, rivals):
    acc = {}
    for c in B.current:
        p = parts[c]
        at31 = calc.daily_normal(B, c, p["cover"], {**p["S"], **{k: v for k, v in p["fc"].items() if v is not None}})(
            dt.date(2026, 12, 31))
        acc[c] = {"name": [x["name"] for x in B.clients if x["key"] == c][0],
                  "on_hand_6dec": int(B.acc[c]["oh"][B.idx[B.last]]), "notices": round(B.notices.get(c, 0.0)),
                  "close_31dec": round(p["oh_dec31"]), "normal_31dec": round(at31),
                  "excess_31dec": round(max(0.0, p["oh_dec31"] - at31)), "burn_share": round(p["s"], 3),
                  "excess_30jun": round(p["e_jun"]),
                  "episodes": [[e[0].isoformat(), round(e[1], 3)] for e in p["episodes"]]}
    seas31 = sum(acc[c]["excess_31dec"] for c in SEASONAL)
    tot31 = sum(acc[c]["excess_31dec"] for c in B.current)
    win = ["%s 2027" % MONTHS[m] for m in range(1, 7)]
    gT, gS = asks.rounded_a2(teu, share)
    jan26 = [B.idx[d] for d in calc.month_days((2026, 1))]
    five_jan = int(sum(B.acc[c]["ship"][jan26].sum() for c in SEASONAL))
    F = {
        "call": {"blocks": g["blocks"], "of": 6, "cost": "$2,200,000", "cost_n": 2200000, "price": "$1,100,000"},
        "components": {
            "on_hand_6dec": int(sum(B.acc[c]["oh"][B.idx[B.last]] for c in B.current)),
            "excess_31dec": int(tot31), "seasonal_excess_31dec": int(seas31),
            "excess_30jun": int(sum(acc[c]["excess_30jun"] for c in B.current)),
            "block_use_pct": [round(100 * u, 1) for u in g["use"]],
            "block_use_whole": [int(round(100 * u)) for u in g["use"]],
            "positions": dict(zip(win, [int(round(x)) for x in g["positions"]])),
            "overflow": dict(zip(win, [int(round(x)) for x in g["overflow"]])),
            "overflow_min": int(round(min(g["overflow"]))), "overflow_max": int(round(max(g["overflow"]))),
            "five_jan2026_shipments": five_jan, "five_excess_half": int(round(seas31 / 2.0)),
        },
        "accounts": acc,
        "steps": {"past_steps": [s.isoformat() for s in B.steps], "notices_total": round(sum(B.notices.values())),
                  "close_31dec_total": round(sum(parts[c]["oh_dec31"] for c in B.current)), "hurdle_pct": 75,
                  "position_months": 24000, "rack_positions": 40000, "block_positions": 4000,
                  "notice_window_days": 90, "policy_sections": ["3.1", "4.2"], "legacy_last_close": "11 September 2022",
                  "window_year": 2027, "legacy_year": 2022, "episode_floor_pct": 12, "episode_min_days": 15,
                  "saturday_weight": "0.40"},
        "rivals": {"fade_half_life_months": round(rivals["fade"]["half_life_months"], 2),
                   "fade_blocks": rivals["fade"]["blocks"],
                   "fade_block1_pct": int(round(100 * rivals["fade"]["use"][0])),
                   "book_burn_blocks": rivals["book_burn"]["blocks"], "book_burn_share": round(rivals["book_burn"]["phi"], 3),
                   "per_account_fall_blocks": rivals["per_account_fall"]["blocks"],
                   "no_payback_blocks": rivals["no_payback"]["blocks"],
                   "book_burn_totals_blocks": rivals["book_burn_totals"]["blocks"],
                   "flat_hold_blocks": flat_hold(B)},
        "a2_teu": {"%s 2026" % MONTHS[k[1]]: gT[k] for k in asks.MONTHS_2026},
        "a2_share": {"%s 2026" % MONTHS[k[1]]: "%.3f" % gS[k] for k in asks.MONTHS_2026},
        "a2_share_raw": {"%s 2026" % MONTHS[k[1]]: share[k] for k in asks.MONTHS_2026},
        "a2_teu_raw": {"%s 2026" % MONTHS[k[1]]: teu[k] for k in asks.MONTHS_2026},
    }
    return F


# ----------------------------------------------------------------------------------- the workbook

def write_workbook(path, B, g, parts, F, inp, teu, share):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DCE6EA")
    call_fill = PatternFill("solid", fgColor="FFF2CC")

    def header(ws, row, names):
        for k, h in enumerate(names, 1):
            c = ws.cell(row=row, column=k, value=h)
            c.font = bold
            c.fill = head_fill
            c.alignment = Alignment(wrap_text=True, vertical="bottom")

    def widths(ws, ws_w):
        for k, w in enumerate(ws_w, 1):
            ws.column_dimensions[get_column_letter(k)].width = w

    # ---- the call
    ws = wb.active
    ws.title = "Blocks"
    ws["A1"] = ("Tamarind Avenue Logistics Park, Building B: the second block would be %.0f%% used over January to "
                "June 2027, the third %.0f%%" % (100 * g["use"][1], 100 * g["use"][2]))
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Capacity planning | forecast from the close of Sunday 6 December 2026 | prepared Monday 7 December 2026"
    ws["A4"] = "Recommendation: take 2 of the 6 blocks (blocks B-1 and B-2), $2,200,000 for the term."
    ws["A4"].font = Font(bold=True, size=12)
    ws["A4"].fill = call_fill
    ws["A5"] = ("Space policy 4.2: a block is taken when the overflow it would hold over the term is at least 75% of "
                "its 24,000 position-months.")
    header(ws, 7, ["Month", "Forecast pallets (mean of daily closes)", "Overflow above 40,000 rack positions",
                   "Held in block 1", "Held in block 2", "Held in block 3"])
    r = 8
    O = g["overflow"]
    for i, m in enumerate(range(1, 7)):
        o = O[i]
        vals = ["%s 2027" % MONTHS[m][:3], int(round(g["positions"][i])), int(round(o)),
                int(round(min(max(o, 0), 4000))), int(round(min(max(o - 4000, 0), 4000))),
                int(round(min(max(o - 8000, 0), 4000)))]
        for k, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=k, value=v)
            if k > 1:
                c.number_format = "#,##0"
        r += 1
    r += 1
    header(ws, r, ["Block", "Positions", "Position-months", "Overflow held over the term (position-months)",
                   "Use of position-months", "Take?"])
    r += 1
    for k in range(6):
        held = float(np.clip(np.asarray(O) - 4000 * k, 0, 4000).sum())
        vals = ["B-%d" % (k + 1), 4000, 24000, int(round(held)), g["use"][k], "Yes" if k < g["blocks"] else "No"]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            if j in (2, 3, 4):
                c.number_format = "#,##0"
            if j == 5:
                c.number_format = "0.0%"
        r += 1
    ws.cell(row=r + 1, column=1, value="Monthly figures are rounded separately; the term totals and uses are computed "
                                       "before rounding. Notes: see the Notes sheet.")
    widths(ws, [10, 20, 22, 16, 16, 16])
    ws.freeze_panes = "A8"

    # ---- by account
    ws = wb.create_sheet("By account")
    seas_e = sum(F["accounts"][c]["excess_31dec"] for c in SEASONAL)
    ws["A1"] = ("Five seasonal accounts hold {:,} of the {:,} pallets over normal stock at 31 December; {:,} are still "
                "here at 30 June".format(seas_e, sum(F["accounts"][c]["excess_31dec"] for c in B.current),
                                         sum(F["accounts"][c]["excess_30jun"] for c in B.current)))
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Each account's excess over its normal stock leaves at the share of its own shipments its excess left at "
                "after the duty steps of 1 September 2019, 5 April 2025 and 7 August 2025, until it is gone.")
    header(ws, 4, ["Client ID", "Client", "On hand, close 6 Dec", "December notices due (pallets)",
                   "Projected close 31 Dec", "Normal stock 31 Dec", "Excess 31 Dec", "Burn share of shipments",
                   "Excess left 30 Jun"])
    r = 5
    for c in B.current:
        a = F["accounts"][c]
        vals = [c, a["name"], a["on_hand_6dec"], a["notices"], a["close_31dec"], a["normal_31dec"], a["excess_31dec"],
                a["burn_share"], a["excess_30jun"]]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=j, value=v)
            if j in (3, 4, 5, 6, 7, 9):
                cell.number_format = "#,##0"
            if j == 8:
                cell.number_format = "0.000"
        r += 1
    tot = ["", "All accounts", sum(F["accounts"][c]["on_hand_6dec"] for c in B.current),
           sum(F["accounts"][c]["notices"] for c in B.current), sum(F["accounts"][c]["close_31dec"] for c in B.current),
           sum(F["accounts"][c]["normal_31dec"] for c in B.current),
           sum(F["accounts"][c]["excess_31dec"] for c in B.current), None,
           sum(F["accounts"][c]["excess_30jun"] for c in B.current)]
    for j, v in enumerate(tot, 1):
        cell = ws.cell(row=r, column=j, value=v)
        cell.font = bold
        if j in (3, 4, 5, 6, 7, 9):
            cell.number_format = "#,##0"
    widths(ws, [10, 30, 14, 16, 14, 14, 12, 14, 14])
    ws.freeze_panes = "C5"

    # ---- A2
    ws = wb.create_sheet("Port TEU 2026")
    ws["A1"] = "Import containers whose freight reached us from the two ports, January to October 2026"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Containers received direct and the containers each Avila transload trailer carried, one count per box "
                "per arrival; TEU per the WMS dictionary; months on Pacific dates. Ports: Los Angeles loaded imports and "
                "Long Beach inbound loaded.")
    header(ws, 4, ["Month", "Our TEU", "Los Angeles loaded imports (TEU)", "Long Beach inbound loaded (TEU)",
                   "Our share of the two ports (%)"])
    r = 5
    for k in asks.MONTHS_2026:
        lab = "%s 2026" % MONTHS[k[1]]
        vals = ["%s 2026" % MONTHS[k[1]][:3], F["a2_teu"][lab], inp["LA"][k], inp["LB"][k], float(F["a2_share"][lab])]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=j, value=v)
            if j in (2, 3, 4):
                cell.number_format = "#,##0"
            if j == 5:
                cell.number_format = "0.000"
        r += 1
    widths(ws, [10, 10, 18, 18, 16])
    ws.freeze_panes = "A5"

    # ---- notes
    ws = wb.create_sheet("Notes")
    notes = [
        ("As of", "The close of Sunday 6 December 2026 (nightly inventory run of 7 December, 08:00Z)."),
        ("Accounts", "The 14 accounts under contract on 7 December. Bellamy Kitchen Co. (left August 2021) and Hollins "
                     "Apparel (left October 2025) hold no stock and are not in the forecast."),
        ("Pallets counted", "Every pallet an account holds with us, in our racks or outside them, the Dominguez "
                            "sublease included (space policy 3.1); the sublease ends on 31 December 2026."),
        ("Month's figure", "The mean of each day's closing stock in the month (space policy 3.1)."),
        ("Overflow", "The month's figure above the 40,000 rack positions of the four buildings (space policy 3.2)."),
        ("Block test", "Blocks fill in order; a block is taken when the overflow it would hold over the six months is at "
                       "least 75% of its 24,000 position-months (space policy 4.2)."),
        ("Normal stock", "Each account's median, over its months from July 2023 clear of the three months before and "
                         "five after each duty step, of pallets on hand over the mean of that and the next month's "
                         "shipments; set at mid-month and interpolated by day."),
        ("December", "Pallets due after the close from the booking tracker; notices in boxes only at the account's own "
                     "pallets per box received over the last 90 days (tracker planning basis). Shipments for the rest "
                     "of December and for 2027: the same month a year earlier, grown at the account's last twelve "
                     "months over the twelve before."),
        ("Burn share", "For each earlier duty step, the fall in the account's excess over normal stock against its "
                       "cumulative shipments from the step (a fitted slope), until the excess fell under a quarter of "
                       "its size or the next advisory; episodes weighted by their shipments."),
        ("Port TEU", "Containers received direct (receipts of one container number within three days are one box; "
                     "the same number months later is a new import) and the containers in each Avila transload "
                     "trailer (Avila's delivery manifest); domestic transfers carry no port container. TEU per the "
                     "WMS export dictionary; months on Pacific dates."),
    ]
    ws["A1"] = "Notes"
    ws["A1"].font = Font(bold=True, size=12)
    for i, (k, v) in enumerate(notes, 3):
        ws.cell(row=i, column=1, value=k).font = bold
        c = ws.cell(row=i, column=2, value=v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, [16, 110])
    wb.save(path)
    PP.normalize_ooxml(path, "Capacity Planning", "Capacity Planning", "2026-12-07T19:02:00Z", "2026-12-07T23:38:00Z",
                       title="Overflow blocks H1 2027", excel=True)


# ----------------------------------------------------------------------------------- the chart

def write_chart(path, F):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    matplotlib.rcParams.update({"font.family": "DejaVu Sans", "pdf.fonttype": 42})
    labs = ["%s 2026" % MONTHS[m] for m in range(1, 11)]
    vals = [float(F["a2_share"][x]) for x in labs]
    teus = [F["a2_teu"][x] for x in labs]
    fig, ax = plt.subplots(figsize=(11, 8.5))
    xs = np.arange(10)
    colours = ["#8fb3bd"] * 8 + ["#1f5f6e"] * 2
    ax.bar(xs, vals, width=0.62, color=colours, zorder=3)
    ax.grid(axis="y", color="#d9dee3", linewidth=0.7, zorder=0)
    for x, v, t in zip(xs, vals, teus):
        ax.text(x, v + 0.004, "{:,} TEU".format(t), ha="center", va="bottom", fontsize=10, color="#1f2328")
        ax.text(x, v - 0.012, "%.3f%%" % v, ha="center", va="top", fontsize=9.5, color="white")
    ax.axvline(7.5, color="#5b626b", linewidth=0.8, linestyle=(0, (4, 3)), zorder=2)
    ax.text(7.42, max(vals) * 1.16, "Duty increase announced\n24 August; front-loading\nfrom September",
            fontsize=9, color="#5b626b", va="top", ha="right")
    ax.set_xticks(xs)
    ax.set_xticklabels([MONTHS[m][:3] for m in range(1, 11)], fontsize=11)
    ax.set_ylabel("Share of the two ports' loaded import TEU (%)", fontsize=11)
    ax.set_ylim(0, max(vals) * 1.18)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FormatStrFormatter("%.2f%%"))
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    ax.set_title("Our share of San Pedro Bay loaded imports reached %.3f%% in September as clients front-loaded"
                 % vals[8], fontsize=13, loc="left", pad=14)
    fig.text(0.07, 0.052, "Bars: our TEU over Los Angeles loaded imports plus Long Beach inbound loaded, January to "
             "October 2026. Labels: TEU of the import containers whose freight reached us", fontsize=8.5,
             color="#5b626b")
    fig.text(0.07, 0.034, "from the two ports, received direct or in Avila transload trailers (each box counted once per "
             "arrival).", fontsize=8.5, color="#5b626b")
    fig.text(0.07, 0.014, "Pellerin Distribution Services, capacity planning, 7 December 2026.", fontsize=8.5,
             color="#5b626b")
    fig.subplots_adjust(left=0.08, right=0.97, top=0.9, bottom=0.12)
    fig.savefig(path, format="pdf", metadata={"Title": "Port share 2026", "Author": "Capacity Planning",
                                              "Creator": "Capacity Planning", "Producer": None,
                                              "CreationDate": dt.datetime(2026, 12, 7, 23, 44, 0,
                                                                          tzinfo=dt.timezone.utc)})
    plt.close(fig)


def write_all(target, gdir, gen):
    B, g, parts, inp, teu, share, rivals = compute(target)
    F = figures(B, g, parts, inp, teu, share, rivals)
    write_workbook(os.path.join(gdir, "overflow_blocks_h1_2027.xlsx"), B, g, parts, F, inp, teu, share)
    write_chart(os.path.join(gdir, "port_share_2026.pdf"), F)
    with open(os.path.join(gen, "figures.json"), "w") as fh:
        json.dump(F, fh, indent=1, sort_keys=True)
        fh.write("\n")
    return {"B": B, "g": g, "parts": parts, "inp": inp, "teu": teu, "share": share, "F": F, "rivals": rivals}
