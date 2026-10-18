"""task161 read-back: the shipped files under target/ read into the calc.Book the golden runs on and into
the ask inputs. Nothing here touches the generator's world: only the bundle.

    import read
    book = read.book(target)              # the main call's inputs
    inp = read.ask_inputs(target)         # A1 and A2 inputs
"""
from __future__ import annotations

import csv
import datetime as dt
import os
import re
from zoneinfo import ZoneInfo

import calc

ONE = dt.timedelta(days=1)
PACIFIC = ZoneInfo("America/Los_Angeles")
F_LEGACY = "wms1_inv_daily_20190401-20220930.txt"
F_NIGHTLY = "wms_inventory_nightly_20220913_20261207.csv"
F_OUTBOUND = "wms_outbound_daily_20220912_20261206.csv"
F_RECEIPTS = "wms_inbound_receipts_20220912_20261206.csv"
F_LPN = "wms_lpn_inventory_20261207T0800Z.csv"
F_CLIENTS = "client_master_2026-12.xlsx"
F_TRACKER = "dec2026_inbound_booking_tracker.xlsx"
F_AVILA = "avila_delivery_manifest_2024-03_2026-12.xlsx"
F_POLA = "pola_container_stats_2026.csv"
F_POLB = "polb_teu_stats_2026.csv"
F_ADVISORIES = "yv_client_advisories_2019-2026.pdf"
F_POLICY = "space_planning_policy_rev4.pdf"
AS_OF = dt.date(2026, 12, 7)
OCEAN = ("22G1", "42G1", "45G1", "L5G1")
MONTH_NO = {m: i for i, m in enumerate(["", "January", "February", "March", "April", "May", "June", "July", "August",
                                        "September", "October", "November", "December"]) if m}


def utc(s):
    return dt.datetime.strptime(s, "%Y-%m-%dT%H:%M:%SZ")


def pacific_date(ts):
    return ts.replace(tzinfo=dt.timezone.utc).astimezone(PACIFIC).date()


def run_close(stamp):
    """A nightly run starts at midnight Pacific and holds the close of the day just ended."""
    return pacific_date(utc(stamp)) - ONE


def pdf_text(path):
    from pypdf import PdfReader
    return "\n".join(p.extract_text() for p in PdfReader(path).pages)


# ----------------------------------------------------------------------------------- the client master

def clients(target):
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(target, F_CLIENTS), data_only=True)
    ws = wb["Clients"]
    rows = list(ws.iter_rows(values_only=True))
    head = list(rows[0])
    out = []
    for r in rows[1:]:
        d = dict(zip(head, r))
        key = d["Client ID"] or d["Legacy code(s)"].split(",")[0].strip()
        end = d["Contract end"]
        out.append({"key": key, "id": d["Client ID"] or None, "name": d["Client name"],
                    "legacy": [x.strip() for x in (d["Legacy code(s)"] or "").split(",") if x.strip()],
                    "start": d["Contract start"].date() if hasattr(d["Contract start"], "date") else d["Contract start"],
                    "end": (end.date() if hasattr(end, "date") else end) if end else None,
                    "line": d["Product line"], "home": d["Home building"]})
    return out


# ----------------------------------------------------------------------------------- the book

def daily_series(target):
    cl = clients(target)
    legacy_map = {code: c["key"] for c in cl for code in c["legacy"]}
    oh, ship, recv = {}, {}, {}
    # legacy to the day before the new system's first close; Sundays carry Saturday's close
    with open(os.path.join(target, F_LEGACY), newline="") as fh:
        for r in csv.DictReader(fh, delimiter="|"):
            d = dt.datetime.strptime(r["BUS_DT"], "%m/%d/%Y").date()
            if d >= dt.date(2022, 9, 12):
                continue
            k = legacy_map[r["CLIENT"]]
            oh[(k, d)] = oh.get((k, d), 0) + int(r["CLOSE_PLT"])
            ship[(k, d)] = ship.get((k, d), 0) + int(r["SHIP_PLT"])
            recv[(k, d)] = recv.get((k, d), 0) + int(r["RCVD_PLT"])
    with open(os.path.join(target, F_NIGHTLY), newline="") as fh:
        for r in csv.DictReader(fh):
            d = run_close(r["snapshot_utc"])
            k = r["client_id"]
            oh[(k, d)] = oh.get((k, d), 0) + int(r["on_hand_pallets"])
    with open(os.path.join(target, F_OUTBOUND), newline="") as fh:
        for r in csv.DictReader(fh):
            d = dt.date.fromisoformat(r["ship_date"])
            ship[(r["client_id"], d)] = ship.get((r["client_id"], d), 0) + int(r["pallets_shipped"])
    with open(os.path.join(target, F_RECEIPTS), newline="") as fh:
        for r in csv.DictReader(fh):
            d = pacific_date(utc(r["received_at_utc"]))
            recv[(r["client_id"], d)] = recv.get((r["client_id"], d), 0) + int(r["pallets"])
    return cl, oh, ship, recv


def advisories(target):
    """The broker's advisory dates and the dates each duty change applied from."""
    txt = re.sub(r"\s+", " ", pdf_text(os.path.join(target, F_ADVISORIES)))
    dated = [dt.datetime.strptime(m, "%d %B %Y").date() for m in re.findall(r"(\d{1,2} [A-Z][a-z]+ \d{4}) \|", txt)]
    steps = [dt.datetime.strptime(m, "%d %B %Y").date()
             for m in re.findall(r"on or after (?:12:01 a\.m\. E[DS]T on )?(\d{1,2} [A-Z][a-z]+ \d{4})", txt)]
    return sorted(set(dated)), sorted(set(steps))


def notices(target, last_close, receipts=None, manifest=None, basis="own"):
    """Pallets due after the close by client: the customer's pallet count, or the boxes at the account's own
    pallets per box received over the last 90 days (ocean boxes by visit and the boxes in transload trailers)."""
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(target, F_TRACKER), data_only=True)
    ws = wb["Dec arrivals"]
    rows = list(ws.iter_rows(min_row=5, values_only=True))
    head = list(rows[0])
    receipts = receipts if receipts is not None else read_receipts(target)
    manifest = manifest if manifest is not None else read_manifest(target)
    ppb = pallets_per_box(receipts, manifest, last_close)
    out = {}
    for r in rows[1:]:
        d = dict(zip(head, r))
        if d["Client ID"] is None:
            continue
        due = d["Due at DC"].date() if hasattr(d["Due at DC"], "date") else d["Due at DC"]
        if due <= last_close:
            continue
        boxes = sum(int(d[h] or 0) for h in ("20'", "40'", "40' HC", "45'"))
        if d["Pallets (customer notice)"] is not None:
            p = float(d["Pallets (customer notice)"])
        elif basis == "zero":
            p = 0.0
        elif basis == "book":
            p = boxes * ppb["__book__"]
        elif basis == "rows":
            p = boxes * ppb["__rows__"][d["Client ID"]]
        else:
            p = boxes * ppb[d["Client ID"]]
        out[d["Client ID"]] = out.get(d["Client ID"], 0.0) + p
    return out


def pallets_per_box(receipts, manifest, last_close, days=90):
    """Pallets per box received over the last `days` days to the close: each ocean box visit's pallets (receipts
    of one container number within three days of each other are one box) and each box a transload trailer carried
    (the partner's manifest); domestic transfers carry no ocean box."""
    lo = last_close - dt.timedelta(days=days - 1)
    src = {}
    for m in manifest:
        src.setdefault(m["trailer"], []).append(m)
    pal, box, rows_p, rows_n = {}, {}, {}, {}
    last = {}
    for r in sorted(receipts, key=lambda x: x["received_at"]):
        d = pacific_date(r["received_at"])
        if not (lo <= d <= last_close):
            continue
        c = r["client"]
        rows_p[c] = rows_p.get(c, 0) + r["pallets"]
        rows_n[c] = rows_n.get(c, 0) + 1
        if r["type"] in OCEAN:
            pal[c] = pal.get(c, 0) + r["pallets"]
            cn = r["equipment"]
            if cn in last and (d - last[cn]).days <= 3:
                last[cn] = d
                continue
            last[cn] = d
            box[c] = box.get(c, 0) + 1
        elif r["equipment"] in src:
            for m in src[r["equipment"]]:
                pal[c] = pal.get(c, 0) + m["pallets"]
                box[c] = box.get(c, 0) + 1
    out = {c: pal[c] / box[c] for c in box}
    out["__book__"] = sum(pal.values()) / sum(box.values())
    out["__rows__"] = {c: rows_p[c] / rows_n[c] for c in rows_n}
    return out


def book(target, notice_basis="own", shift_run_day=0, stack_parallel=False, sundays="carry", legacy_codes="all",
         include_former=True):
    """The Book the golden runs on. The keyword arguments are the Layer 1 mishandlings checks.py runs
    (R2): the notices' box-only rows, the run stamp read as the close date, the parallel-run days stacked
    from both systems, the legacy Sundays left out of the month's mean, a second legacy code unmapped."""
    cl, oh, ship, recv = daily_series(target)
    if legacy_codes != "all" or stack_parallel or shift_run_day:
        cl, oh, ship, recv = _mishandled(target, legacy_codes, stack_parallel, shift_run_day)
    last_close = max(d for (_, d) in oh)
    first = min(d for (_, d) in oh)
    days = []
    d = first
    while d <= last_close:
        days.append(d)
        d += ONE
    import numpy as np
    acc = {}
    for c in cl:
        k = c["key"]
        if not include_former and c["end"] is not None:
            continue
        o = np.zeros(len(days))
        s = np.zeros(len(days))
        rc = np.zeros(len(days))
        prev = 0.0
        for i, dd in enumerate(days):
            if (k, dd) in oh:
                prev = float(oh[(k, dd)])
            elif c["end"] is not None and dd > c["end"]:
                prev = 0.0
            elif dd < c["start"]:
                prev = 0.0
            o[i] = prev
            s[i] = ship.get((k, dd), 0)
            rc[i] = recv.get((k, dd), 0)
        acc[k] = {"oh": o, "ship": s, "recv": rc}
    current = [c["key"] for c in cl if c["end"] is None or c["end"] >= AS_OF]
    announced, steps = advisories(target)
    past = [x for x in steps if x <= last_close]
    nts = notices(target, last_close, basis=notice_basis)
    B = calc.Book(days, acc, nts, past, current, last_close, announced=announced)
    if sundays == "rows":
        # the legacy months' mean taken over the rows on file (Monday to Saturday)
        B.skip_day = lambda x: x < dt.date(2022, 9, 12) and x.weekday() == 6
    B.clients = cl
    B.future_steps = [x for x in steps if x > last_close]
    return B


def _mishandled(target, legacy_codes, stack_parallel, shift_run_day):
    cl = clients(target)
    legacy_map = {}
    for c in cl:
        codes = c["legacy"] if legacy_codes == "all" else c["legacy"][:1]
        for code in codes:
            legacy_map[code] = c["key"]
    oh, ship, recv = {}, {}, {}
    cut = dt.date(2022, 10, 1) if stack_parallel else dt.date(2022, 9, 12)
    with open(os.path.join(target, F_LEGACY), newline="") as fh:
        for r in csv.DictReader(fh, delimiter="|"):
            d = dt.datetime.strptime(r["BUS_DT"], "%m/%d/%Y").date()
            if d >= cut or r["CLIENT"] not in legacy_map:
                continue
            k = legacy_map[r["CLIENT"]]
            oh[(k, d)] = oh.get((k, d), 0) + int(r["CLOSE_PLT"])
            ship[(k, d)] = ship.get((k, d), 0) + int(r["SHIP_PLT"])
            recv[(k, d)] = recv.get((k, d), 0) + int(r["RCVD_PLT"])
    with open(os.path.join(target, F_NIGHTLY), newline="") as fh:
        for r in csv.DictReader(fh):
            d = run_close(r["snapshot_utc"]) + dt.timedelta(days=shift_run_day)
            k = r["client_id"]
            oh[(k, d)] = oh.get((k, d), 0) + int(r["on_hand_pallets"])
    with open(os.path.join(target, F_OUTBOUND), newline="") as fh:
        for r in csv.DictReader(fh):
            d = dt.date.fromisoformat(r["ship_date"])
            ship[(r["client_id"], d)] = ship.get((r["client_id"], d), 0) + int(r["pallets_shipped"])
    with open(os.path.join(target, F_RECEIPTS), newline="") as fh:
        for r in csv.DictReader(fh):
            d = pacific_date(utc(r["received_at_utc"]))
            recv[(r["client_id"], d)] = recv.get((r["client_id"], d), 0) + int(r["pallets"])
    if shift_run_day:
        last = max(d for (_, d) in oh)
        oh = {kd: v for kd, v in oh.items() if kd[1] <= last}
    return cl, oh, ship, recv


# ----------------------------------------------------------------------------------- the asks' inputs

def read_receipts(target):
    out = []
    with open(os.path.join(target, F_RECEIPTS), newline="") as fh:
        for r in csv.DictReader(fh):
            out.append({"received_at": utc(r["received_at_utc"]), "type": r["equipment_type"],
                        "equipment": r["equipment_id"], "pallets": int(r["pallets"]), "client": r["client_id"],
                        "site": r["site_id"]})
    return out


def read_manifest(target):
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(target, F_AVILA), read_only=True, data_only=True)
    ws = wb["Deliveries"]
    rows = ws.iter_rows(min_row=4, values_only=True)
    head = list(next(rows))
    out = []
    for r in rows:
        d = dict(zip(head, r))
        if d.get("Trailer") is None:
            continue
        out.append({"trailer": d["Trailer"], "container": d["Container"], "iso": d["ISO Type"],
                    "pallets": int(d["Pallets"]), "port": d["Port"], "terminal": d["Pickup Terminal"],
                    "delivered": d["Delivery Date"].date() if hasattr(d["Delivery Date"], "date") else d["Delivery Date"]})
    return out


def read_ports(target):
    LA, LB = {}, {}
    with open(os.path.join(target, F_POLA), newline="") as fh:
        for r in csv.DictReader(fh):
            parts = r["Month"].split()
            if len(parts) == 2 and parts[0] in MONTH_NO:
                LA[(int(parts[1]), MONTH_NO[parts[0]])] = int(r["Loaded Imports"])
    ab = {m[:3].upper(): i for m, i in MONTH_NO.items()}
    with open(os.path.join(target, F_POLB), newline="") as fh:
        for r in csv.DictReader(fh):
            LB[(int(r["Year"]), ab[r["Month"]])] = int(r["Inbound Loaded"])
    return LA, LB


def site_onhand(target, close):
    out = {}
    with open(os.path.join(target, F_NIGHTLY), newline="") as fh:
        for r in csv.DictReader(fh):
            if run_close(r["snapshot_utc"]) == close:
                out[r["site_id"]] = out.get(r["site_id"], 0) + int(r["on_hand_pallets"])
    return out


def read_lpn(target):
    out = []
    with open(os.path.join(target, F_LPN), newline="") as fh:
        for r in csv.DictReader(fh):
            out.append({"site": r["site_id"], "location_type": r["location_type"], "lpn": r["lpn"],
                        "client": r["client_id"], "snapshot_utc": r["snapshot_utc"]})
    return out


def ask_inputs(target):
    LA, LB = read_ports(target)
    return {"receipts": read_receipts(target), "manifest": read_manifest(target), "LA": LA, "LB": LB,
            "to_local": pacific_date}
