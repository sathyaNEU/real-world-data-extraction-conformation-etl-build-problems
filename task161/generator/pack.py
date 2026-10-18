"""task161 pack: every file under target/, written from the built world (core.build_world).

The folder is the capacity planner's on Monday 7 December 2026: the legacy WMS1 daily inventory to its
retirement, the new WMS's nightly inventory, outbound and inbound exports and its pallet-level export
at the close of 6 December, the client master and the site list, customer service's December booking
tracker, the transload partner's delivery manifest, the two ports' monthly container statistics, the
space planning policy, the park's block offer, the customs broker's advisories, the export dictionary
and the export manifest, and two working files the solution never reads (the fiscal-period capacity
report and the H1 2027 dock labour plan).
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import os

import numpy as np

import docs as D
import postprocess as PP
import records
import sim
import world as W

ONE = dt.timedelta(days=1)

F_LEGACY = "wms1_inv_daily_20190401-20220930.txt"
F_NIGHTLY = "wms_inventory_nightly_20220913_20261207.csv"
F_OUTBOUND = "wms_outbound_daily_20220912_20261206.csv"
F_RECEIPTS = "wms_inbound_receipts_20220912_20261206.csv"
F_LPN = "wms_lpn_inventory_20261207T0800Z.csv"
F_CLIENTS = "client_master_2026-12.xlsx"
F_SITES = "site_list.csv"
F_TRACKER = "dec2026_inbound_booking_tracker.xlsx"
F_AVILA = "avila_delivery_manifest_2024-03_2026-12.xlsx"
F_POLA = "pola_container_stats_2026.csv"
F_POLB = "polb_teu_stats_2026.csv"
F_POLICY = "space_planning_policy_rev4.pdf"
F_OFFER = "ridley-ames_bldg-b_overflow_blocks_2026-11-19.pdf"
F_ADVISORIES = "yv_client_advisories_2019-2026.pdf"
F_DICT = "wms_export_dictionary_v3.pdf"
F_MANIFEST = "export_manifest.csv"
F_CAPACITY = "capacity_report_fy2026_p12.xlsx"
F_LABOR = "dock_labor_plan_h1_2027.xlsx"
DISTRACTORS = [F_CAPACITY, F_LABOR]
ALL_FILES = [F_LEGACY, F_NIGHTLY, F_OUTBOUND, F_RECEIPTS, F_LPN, F_CLIENTS, F_SITES, F_TRACKER, F_AVILA, F_POLA,
             F_POLB, F_POLICY, F_OFFER, F_ADVISORIES, F_DICT, F_MANIFEST, F_CAPACITY, F_LABOR]

SITES = [b[0] for b in W.BUILDINGS]
PRODUCT_LINE = {"furniture": "Furniture", "housewares": "Housewares", "toys": "Toys",
                "electronics": "Mobile and electronics accessories", "apparel": "Apparel", "footwear": "Footwear",
                "textiles": "Home textiles", "outdoor": "Patio and garden", "pet": "Pet supplies",
                "auto": "Auto parts and accessories", "decor": "Seasonal decor", "appliances": "Small appliances",
                "fitness": "Fitness equipment", "office": "Office products"}
MONTHS = ["", "January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
          "November", "December"]


def contract(code):
    v = W.ACCOUNTS[code]
    return dt.date.fromisoformat(v[6]), (dt.date.fromisoformat(v[7]) if v[7] else None)


def active(code, d):
    a, b = contract(code)
    return d >= a and (b is None or d <= b)


def zstamp(ts):
    return ts.strftime("%Y-%m-%dT%H:%M:%SZ")


def write_csv(path, header, rows, crlf=False, delim=","):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n" if crlf else "\n", delimiter=delim)
    w.writerow(header)
    n = 0
    for r in rows:
        w.writerow(r)
        n += 1
    with open(path, "w", encoding="utf-8", newline="") as fh:
        fh.write(buf.getvalue())
    return n


# ----------------------------------------------------------------------------------- the legacy WMS1 export

def legacy_rows(T):
    """One row per legacy client code per business day (Monday to Saturday) to the close of 30 September
    2022. Rosales ran two legacy codes (basics and kids); its flows and stock are split between them."""
    rng = np.random.default_rng([W.SEED, 7101])
    days = [d for d in sim.ALL_DAYS if d <= W.OLD_WMS_END and d.weekday() != 6]
    out = []
    for c in sorted(W.ACCOUNTS):
        codes = W.ACCOUNT_NAMES[c][2]
        if not codes:
            continue
        S, R, OH = T[c]["S"], T[c]["R"], T[c]["OH"]
        b_close, share = None, W.ROSALES_SPLIT
        for d in days:
            if not active(c, d):
                continue
            i = sim.DIDX[d]
            close, rcv, shp = int(OH[i]), int(R[i]), int(S[i])
            opn = close - rcv + shp
            if active(c, d - ONE):
                j = sim.DIDX[d - ONE]
                assert opn == int(OH[j]), (c, d)
            if len(codes) == 1:
                out.append((d, codes[0], opn, rcv, shp, close))
                continue
            # two legacy codes: the basics share drifts slowly around its mean
            share = W.ROSALES_SPLIT + 0.92 * (share - W.ROSALES_SPLIT) + float(rng.normal(0, 0.006))
            if b_close is None:
                b_close = int(round(opn * share))
            b_open = b_close
            rb = int(round(rcv * float(np.clip(share + rng.normal(0, 0.05), 0.4, 0.85))))
            target = int(round(close * share))
            lo, hi = max(0, b_open + rb - close), min(shp, b_open + rb)
            sb = int(min(max(b_open + rb - target, lo), hi))
            b_close = b_open + rb - sb
            out.append((d, codes[0], b_open, rb, sb, b_close))
            out.append((d, codes[1], opn - b_open, rcv - rb, shp - sb, close - b_close))
    out.sort(key=lambda r: (r[0], r[1]))
    return [(d.strftime("%m/%d/%Y"), code, o, r, s, c) for d, code, o, r, s, c in out]


# ----------------------------------------------------------------------------------- the new WMS exports

def nightly_rows(T):
    """The nightly inventory run: every client under contract by site, at every close from 12 September
    2022 to 6 December 2026, the sublease as a site while it was let."""
    out = []
    splits = {}
    d = W.NEW_WMS_START
    while d <= W.LAST_CLOSE:
        split = records.site_split(T, d)
        splits[d] = split
        sites = SITES + (["DHS1"] if records.sublease_open(d) else [])
        stamp = zstamp(records.run_stamp(d))
        for c in sorted(W.ACCOUNTS, key=lambda x: W.ACCOUNT_NAMES[x][1] or ""):
            pid = W.ACCOUNT_NAMES[c][1]
            if pid is None or not active(c, d):
                continue
            tot = records.site_totals(split[c]) if c in split else {}
            i = sim.DIDX[d]
            assert sum(tot.values()) == int(T[c]["OH"][i]), (c, d)
            assert all(tot.get(s, 0) == 0 for s in tot if s not in sites), (c, d)
            for s in sites:
                out.append((stamp, s, pid, int(tot.get(s, 0))))
        d += ONE
    out.sort(key=lambda r: (r[0], r[1], r[2]))
    return out, splits


def outbound_rows(T):
    rng = np.random.default_rng([W.SEED, 7102])
    out = []
    d = W.NEW_WMS_START
    while d <= W.LAST_CLOSE:
        i = sim.DIDX[d]
        for c in sorted(W.ACCOUNTS, key=lambda x: W.ACCOUNT_NAMES[x][1] or ""):
            pid = W.ACCOUNT_NAMES[c][1]
            if pid is None:
                continue
            p = int(T[c]["S"][i])
            if p <= 0:
                continue
            loads = max(1, int(np.ceil(p / float(rng.uniform(16.0, 24.0)))))
            out.append((d.isoformat(), pid, loads, p))
        d += ONE
    return out


def receipt_rows(receipts):
    rng = np.random.default_rng([W.SEED, 7103])
    cap = np.array([b[4] for b in W.BUILDINGS], float)
    rows = [r for r in receipts if r["day"] >= W.NEW_WMS_START]
    rows.sort(key=lambda r: (r["received_at"], r["account"], r["equipment"], r.get("part", 0)))
    out = []
    no = 4180000
    for r in rows:
        no += int(rng.integers(1, 3))
        home = W.HOME[r["account"]]
        if rng.random() < 0.86:
            site = home
        else:
            site = SITES[int(rng.choice(4, p=cap / cap.sum()))]
        out.append(("RC%08d" % no, zstamp(r["received_at"]), site, W.ACCOUNT_NAMES[r["account"]][1],
                    r["equipment"], r["type"], int(r["pallets"])))
    return out


def lpn_rows(lpn):
    stamp = zstamp(records.run_stamp(W.LAST_CLOSE))
    rows = sorted(lpn, key=lambda r: (r["site"], r["location"], r["lpn"]))
    return [(stamp, r["site"], r["location"], r["location_type"], r["account"], r["lpn"], r["item"], r["qty_cases"],
             r["received"].isoformat()) for r in rows]


# ----------------------------------------------------------------------------------- workbooks

def _wb():
    import openpyxl
    return openpyxl.Workbook()


def _bold(cell, size=None):
    from openpyxl.styles import Font
    cell.font = Font(bold=True, size=size) if size else Font(bold=True)


def _widths(ws, widths):
    from openpyxl.utils import get_column_letter
    for k, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(k)].width = w


def _finish(path, wb, creator, last_by, created, modified, title, company=None):
    wb.save(path)
    PP.normalize_ooxml(path, creator, last_by, created, modified, title=title, excel=True, company=company)


def write_clients(path):
    wb = _wb()
    ws = wb.active
    ws.title = "Clients"
    head = ["Client ID", "Client name", "Legacy code(s)", "Product line", "Contract start", "Contract end",
            "Home building"]
    for k, h in enumerate(head, 1):
        _bold(ws.cell(row=1, column=k, value=h))
    r = 2
    for c in sorted(W.ACCOUNTS, key=lambda x: (W.ACCOUNT_NAMES[x][1] is None, W.ACCOUNT_NAMES[x][1] or x)):
        name, pid, legacy = W.ACCOUNT_NAMES[c]
        a, b = contract(c)
        vals = [pid or "", name, ", ".join(legacy), PRODUCT_LINE[W.ACCOUNTS[c][0]], a, b, W.HOME[c]]
        for k, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=k, value=v)
            if isinstance(v, dt.date):
                cell.number_format = "yyyy-mm-dd"
        r += 1
    ws.freeze_panes = "A2"
    _widths(ws, [10, 30, 20, 32, 14, 14, 14])
    _finish(path, wb, "Jerome Lam", "Nicole Freeman", "2022-08-22T17:04:00Z", "2026-11-30T23:41:00Z", "Client master")


def site_rows():
    yard = {"CAR1": 60, "RDZ1": 52, "WIL1": 36, "CMP1": 30}
    doors = {"CAR1": 38, "RDZ1": 34, "WIL1": 26, "CMP1": 22}
    out = []
    for code, legacy, name, city, cap in W.BUILDINGS:
        out.append((code, legacy, name, city, "CA", cap, doors[code], yard[code], "Owned building", "", ""))
    a, b = records.SUBLEASE_STINTS[-1]
    out.append(("DHS1", "", W.SUBLEASE_SITE[1].title(), W.SUBLEASE_SITE[2], "CA", records.DHS_CAP, 12, 0,
                "Sublease (also let 2025-07-14 to 2025-10-31)", a.isoformat(), b.isoformat()))
    return out


def write_tracker(path, notices):
    wb = _wb()
    ws = wb.active
    ws.title = "Dec arrivals"
    ws["A1"] = "December 2026 inbound: booking notices for arrivals due after the close of Sunday 6 December"
    _bold(ws["A1"], 12)
    ws["A2"] = ("Planning basis: where a customer's notice gives boxes only, plan pallets at that account's own "
                "pallets per box received over the last 90 days.")
    ws["A3"] = "Updated Sunday 6 December 2026, 18:40. Customer service."
    head = ["Client", "Client ID", "Booking", "Carrier", "Vessel", "Terminal", "ETA", "Due at DC", "20'", "40'",
            "40' HC", "45'", "Pallets (customer notice)"]
    for k, h in enumerate(head, 1):
        _bold(ws.cell(row=5, column=k, value=h))
    r = 6
    for n in sorted(notices, key=lambda x: (x["receive_by"], W.ACCOUNT_NAMES[x["account"]][1], x["booking"])):
        name, pid, _ = W.ACCOUNT_NAMES[n["account"]]
        vals = [name, pid, n["booking"], n["carrier"], n["vessel"], n["terminal"], n["eta"], n["receive_by"],
                n["boxes"]["20"] or None, n["boxes"]["40"] or None, n["boxes"]["40HC"] or None,
                n["boxes"]["45"] or None, n["pallets"]]
        for k, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=k, value=v)
            if isinstance(v, dt.date):
                cell.number_format = "d-mmm"
        r += 1
    ws.freeze_panes = "A6"
    _widths(ws, [30, 9, 13, 10, 24, 24, 8, 10, 5, 5, 6, 5, 12])
    _finish(path, wb, "Nicole Freeman", "Nicole Freeman", "2026-11-23T16:12:00Z", "2026-12-07T02:40:00Z",
            "December 2026 inbound booking tracker")


def write_avila(path, manifest):
    rng = np.random.default_rng([W.SEED, 7104])
    wb = _wb()
    ws = wb.active
    ws.title = "Deliveries"
    head = ["Delivery Date", "Trailer", "Container", "ISO Type", "Pallets", "Consignee", "Pickup Terminal"]
    for k, h in enumerate(head, 1):
        _bold(ws.cell(row=1, column=k, value=h))
    rows = sorted(manifest, key=lambda m: (m["delivered"], m["trailer"], m["container"]))
    term = {}
    r = 2
    for m in rows:
        if m["trailer"] not in term:
            term[m["trailer"]] = records.TERMINALS[int(rng.integers(0, len(records.TERMINALS)))]
        vals = [m["delivered"], m["trailer"], m["container"], m["iso"], int(m["pallets"]),
                "Pellerin Distribution / " + W.ACCOUNT_NAMES[m["account"]][0], term[m["trailer"]]]
        for k, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=k, value=v)
            if isinstance(v, dt.date):
                cell.number_format = "mm/dd/yyyy"
        r += 1
    ws.freeze_panes = "A2"
    _widths(ws, [13, 12, 14, 9, 8, 52, 24])
    _finish(path, wb, "Avila Dispatch", "Avila Dispatch", "2024-03-04T15:20:00Z", "2026-12-07T01:15:00Z",
            "Delivery manifest", company=W.TRANSLOAD)


# ----------------------------------------------------------------------------------- the ports

def port_rows(LA, LB):
    rng = np.random.default_rng([W.SEED, 7105])
    la, lb = [], []
    for k in sorted(LA):
        y, m = k
        li = LA[k]
        ei = int(li * float(rng.uniform(0.008, 0.024)))
        lx = int(float(rng.uniform(102000, 131000)))
        ex = int(li * float(rng.uniform(0.62, 0.74)))
        la.append(("%s %d" % (MONTHS[m], y), li, ei, li + ei, lx, ex, lx + ex, li + ei + lx + ex))
        bi = LB[k]
        be = int(bi * float(rng.uniform(0.03, 0.07)))
        bx = int(float(rng.uniform(108000, 129000)))
        bo = int(bi * float(rng.uniform(0.60, 0.70)))
        lb.append((y, MONTHS[m][:3].upper(), bi, be, bx, bo, bi + be + bx + bo))
    tot = [sum(r[j] for r in la) for j in range(1, 8)]
    la.append(("Calendar YTD 2026",) + tuple(tot))
    return la, lb


# ----------------------------------------------------------------------------------- distractors

def fiscal_periods():
    """Fiscal 2026 runs December 2025 to November 2026 in calendar-month periods."""
    out = []
    y, m = 2025, 12
    for p in range(1, 13):
        a = dt.date(y, m, 1)
        b = dt.date(y + (m == 12), m % 12 + 1, 1) - ONE
        out.append(("P%d" % p, a, b))
        y, m = (y + (m == 12), m % 12 + 1)
    return out


def write_capacity(path, nightly):
    """Fiscal 2026 to period 12: mean and peak pallets on hand by site from the nightly run, against rack
    positions. Built from the shipped nightly rows."""
    from openpyxl.styles import Font
    by = {}
    for stamp, site, pid, p in nightly:
        ts = dt.datetime.strptime(stamp, "%Y-%m-%dT%H:%M:%SZ")
        close = records.utc_to_pacific_date(ts) - ONE
        by.setdefault((close, site), 0)
        by[(close, site)] += p
    caps = {b[0]: b[4] for b in W.BUILDINGS}
    caps["DHS1"] = records.DHS_CAP
    wb = _wb()
    ws = wb.active
    ws.title = "FY2026"
    ws["A1"] = "Capacity report, fiscal 2026, periods 1 to 12 (December 2025 to November 2026)"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Pallets on hand by site: mean and peak of the nightly inventory run's closing balances in the period; "
                "utilisation is the mean against the site's rack positions.")
    ws["A3"] = "Capacity planning, 2 December 2026."
    head = ["Period", "From", "To", "Site", "Rack positions", "Mean pallets on hand", "Peak pallets on hand",
            "Utilisation"]
    for k, h in enumerate(head, 1):
        ws.cell(row=5, column=k, value=h).font = Font(bold=True)
    r = 6
    rows = []
    for p, a, b in fiscal_periods():
        for s in SITES + ["DHS1"]:
            vals = []
            d = a
            while d <= b:
                if (d, s) in by:
                    vals.append(by[(d, s)])
                d += ONE
            if not vals:
                continue
            mean = sum(vals) / len(vals)
            rows.append((p, a, b, s, caps[s], round(mean), max(vals), round(mean / caps[s], 4)))
    for row in rows:
        for k, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=k, value=v)
            if isinstance(v, dt.date):
                cell.number_format = "yyyy-mm-dd"
            if k == 8:
                cell.number_format = "0.0%"
            if k in (5, 6, 7):
                cell.number_format = "#,##0"
        r += 1
    ws.freeze_panes = "A6"
    _widths(ws, [8, 12, 12, 7, 14, 20, 20, 12])
    _finish(path, wb, "Capacity Planning", "Capacity Planning", "2025-12-03T18:30:00Z", "2026-12-02T19:05:00Z",
            "Capacity report FY2026")
    return rows


def write_labor(path, T, receipts):
    """The H1 2027 dock labour plan by building: handled pallets per week planned at the same month of 2026
    plus 4 per cent, receiving and shipping shifts at 1,150 pallets handled per shift."""
    from openpyxl.styles import Font
    recv = {}
    for r in receipts:
        if r["day"].year == 2026 and r["day"].month <= 6:
            k = (W.HOME[r["account"]], r["day"].month)
            recv[k] = recv.get(k, 0) + r["pallets"]
    ship = {}
    for c in W.CURRENT:
        for m in range(1, 7):
            a = dt.date(2026, m, 1)
            b = dt.date(2026, m + 1, 1) - ONE
            tot = int(T[c]["S"][sim.DIDX[a]:sim.DIDX[b] + 1].sum())
            ship[(W.HOME[c], m)] = ship.get((W.HOME[c], m), 0) + tot
    wb = _wb()
    ws = wb.active
    ws.title = "H1 2027"
    ws["A1"] = "Dock labor plan, January to June 2027"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Basis: pallets received and shipped by each building's home accounts in the same month of 2026, plus "
                "4 per cent; one receiving or shipping shift per 1,150 pallets handled in a week; 9 FTE per shift.")
    ws["A3"] = "Operations, draft of 30 November 2026."
    head = ["Building", "Month", "Planned received per week", "Planned shipped per week", "Receiving shifts per week",
            "Shipping shifts per week", "Planned FTE"]
    for k, h in enumerate(head, 1):
        ws.cell(row=5, column=k, value=h).font = Font(bold=True)
    r = 6
    rows = []
    for s in SITES:
        for m in range(1, 7):
            weeks = ((dt.date(2026, m + 1, 1) if m < 12 else dt.date(2027, 1, 1)) - dt.date(2026, m, 1)).days / 7.0
            rw = recv.get((s, m), 0) * 1.04 / weeks
            sw = ship.get((s, m), 0) * 1.04 / weeks
            rs = max(1, int(np.ceil(rw / 1150.0)))
            ss = max(1, int(np.ceil(sw / 1150.0)))
            rows.append((s, "%s 2027" % MONTHS[m][:3], int(round(rw)), int(round(sw)), rs, ss, 9 * (rs + ss)))
    for row in rows:
        for k, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=k, value=v)
            if k in (3, 4):
                cell.number_format = "#,##0"
        r += 1
    _widths(ws, [10, 10, 24, 24, 24, 22, 12])
    _finish(path, wb, "Mariah Norman", "Mariah Norman", "2026-11-24T17:50:00Z", "2026-11-30T22:16:00Z",
            "Dock labor plan H1 2027")
    return rows


# ----------------------------------------------------------------------------------- the export manifest

def manifest_rows(counts):
    pulled = "2026-12-07T16:05:00Z"
    return [
        (F_LEGACY, "WMS1 (legacy)", "daily inventory by client: opening, received, shipped, closing pallets",
         "2019-04-01 to 2022-09-30", "2022-10-03T15:40:00Z", counts[F_LEGACY]),
        (F_NIGHTLY, "WMS", "nightly inventory run: pallets on hand by site and client", "closes of 2022-09-12 to 2026-12-06",
         pulled, counts[F_NIGHTLY]),
        (F_OUTBOUND, "WMS", "outbound loads and pallets shipped by client and day", "2022-09-12 to 2026-12-06", pulled,
         counts[F_OUTBOUND]),
        (F_RECEIPTS, "WMS", "inbound receipts with equipment", "2022-09-12 to 2026-12-06", pulled, counts[F_RECEIPTS]),
        (F_LPN, "WMS", "pallet-level inventory by warehouse location", "close of 2026-12-06", pulled, counts[F_LPN]),
        (F_CLIENTS, "WMS client master", "clients, legacy codes, contracts, home buildings", "as of 2026-12-07", pulled,
         counts[F_CLIENTS]),
        (F_SITES, "Facilities", "sites, rack positions, dock doors, yard spots", "as of 2026-12-07", pulled,
         counts[F_SITES]),
        (F_TRACKER, "Customer service", "booking notices for December arrivals", "arrivals due 2026-12-07 to 2026-12-31",
         "2026-12-07T02:40:00Z", counts[F_TRACKER]),
        (F_AVILA, "Avila Transload Services", "trailer deliveries with the ocean containers each carried",
         "2024-03-04 to 2026-12-06", "2026-12-07T01:15:00Z", counts[F_AVILA]),
        (F_POLA, "Port of Los Angeles website", "container statistics (TEU), monthly", "January to October 2026",
         "2026-12-04T19:22:00Z", counts[F_POLA]),
        (F_POLB, "Port of Long Beach website", "TEU statistics, monthly", "January to October 2026",
         "2026-12-04T19:26:00Z", counts[F_POLB]),
        (F_POLICY, "Operations", "space planning policy, revision 4", "effective 2026-03-02", "", ""),
        (F_OFFER, "Ridley-Ames Industrial Properties", "block offer, Tamarind Avenue Logistics Park Building B",
         "dated 2026-11-19", "", ""),
        (F_ADVISORIES, "Yamamoto & Vance Customs Brokers", "client advisories on additional duties, 2019 to 2026",
         "printed 2026-12-03", "", ""),
        (F_DICT, "Information Systems", "WMS exports: field dictionary, version 3", "dated 2026-11-02", "", ""),
        (F_CAPACITY, "Capacity planning", "capacity report, fiscal 2026 periods 1 to 12", "2025-12-01 to 2026-11-30",
         "2026-12-02T19:05:00Z", counts[F_CAPACITY]),
        (F_LABOR, "Operations", "dock labor plan, January to June 2027 (draft)", "2027-01 to 2027-06",
         "2026-11-30T22:16:00Z", counts[F_LABOR]),
    ]


# ----------------------------------------------------------------------------------- everything

def write_all(out, Wd):
    T = Wd["truth"]
    counts = {}
    counts[F_LEGACY] = write_csv(os.path.join(out, F_LEGACY),
                                 ["BUS_DT", "CLIENT", "OPEN_PLT", "RCVD_PLT", "SHIP_PLT", "CLOSE_PLT"],
                                 legacy_rows(T), crlf=True, delim="|")
    nightly, splits = nightly_rows(T)
    counts[F_NIGHTLY] = write_csv(os.path.join(out, F_NIGHTLY), ["snapshot_utc", "site_id", "client_id",
                                                                  "on_hand_pallets"], nightly)
    counts[F_OUTBOUND] = write_csv(os.path.join(out, F_OUTBOUND), ["ship_date", "client_id", "loads",
                                                                    "pallets_shipped"], outbound_rows(T))
    counts[F_RECEIPTS] = write_csv(os.path.join(out, F_RECEIPTS),
                                   ["receipt_no", "received_at_utc", "site_id", "client_id", "equipment_id",
                                    "equipment_type", "pallets"], receipt_rows(Wd["receipts"]))
    counts[F_LPN] = write_csv(os.path.join(out, F_LPN),
                              ["snapshot_utc", "site_id", "location_id", "location_type", "client_id", "lpn", "item",
                               "cases", "received_date"], lpn_rows(Wd["lpn"]))
    write_clients(os.path.join(out, F_CLIENTS))
    counts[F_CLIENTS] = len(W.ACCOUNTS)
    counts[F_SITES] = write_csv(os.path.join(out, F_SITES),
                                ["site_id", "legacy_whse", "site_name", "city", "state", "rack_positions", "dock_doors",
                                 "yard_trailer_spots", "status", "held_from", "held_to"], site_rows())
    write_tracker(os.path.join(out, F_TRACKER), Wd["notices"])
    counts[F_TRACKER] = len(Wd["notices"])
    write_avila(os.path.join(out, F_AVILA), Wd["manifest"])
    counts[F_AVILA] = len(Wd["manifest"])
    la, lb = port_rows(Wd["LA"], Wd["LB"])
    counts[F_POLA] = write_csv(os.path.join(out, F_POLA), ["Month", "Loaded Imports", "Empty Imports", "Total Imports",
                                                            "Loaded Exports", "Empty Exports", "Total Exports",
                                                            "Total TEUs"], la)
    counts[F_POLB] = write_csv(os.path.join(out, F_POLB), ["Year", "Month", "Inbound Loaded", "Inbound Empty",
                                                            "Outbound Loaded", "Outbound Empty", "Total"], lb)
    D.write_policy(os.path.join(out, F_POLICY))
    D.write_offer(os.path.join(out, F_OFFER))
    D.write_advisories(os.path.join(out, F_ADVISORIES))
    D.write_dictionary(os.path.join(out, F_DICT))
    counts[F_CAPACITY] = len(write_capacity(os.path.join(out, F_CAPACITY), nightly))
    counts[F_LABOR] = len(write_labor(os.path.join(out, F_LABOR), T, Wd["receipts"]))
    write_csv(os.path.join(out, F_MANIFEST), ["file", "source", "contents", "covers", "pulled_utc", "rows"],
              manifest_rows(counts))
    counts[F_MANIFEST] = len(ALL_FILES) - 1
    return counts
