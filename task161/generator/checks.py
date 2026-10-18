"""task161 checks: every assertion the build stands on, run on the cut pack (target/) and on the generator's
clean truth. Labels: R1 to R3c the client's three rules executed (dataset-generation 7.1), L the ladder and
the call's margins, SF1 to SF12 the semantic fork table, A the asks, G the input gates, S the realism sweep,
P the pins and the counter-pin sweep, F the figures.

    python3 checks.py            # builds the world and reads target/ (run build.py first)
"""
from __future__ import annotations

import csv
import datetime as dt
import itertools
import json
import os
import re
import shutil
import sys
import tempfile
import warnings
import zipfile

warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import numpy as np   # noqa: E402

import asks          # noqa: E402
import calc          # noqa: E402
import cracker       # noqa: E402
import docs as D     # noqa: E402
import pack as PK    # noqa: E402
import read          # noqa: E402
import records       # noqa: E402
import sim           # noqa: E402
import world as W    # noqa: E402

AS_OF = dt.date(2026, 12, 7)
SEASONAL = ["PD2047", "PD2052", "PD2058", "PD2102", "PD2108"]
PID = {c: W.ACCOUNT_NAMES[c][1] for c in W.ACCOUNTS}


class Sheet:
    def __init__(self):
        self.rows = []

    def ok(self, label, cond, note=""):
        self.rows.append((label, bool(cond), note))
        return cond

    def report(self):
        bad = [r for r in self.rows if not r[1]]
        for label, good, note in self.rows:
            print("  %s  %-6s %s" % ("PASS" if good else "FAIL", label, note))
        print("checks: %d assertions, %d failed" % (len(self.rows), len(bad)))
        if bad:
            raise SystemExit("checks failed: " + ", ".join(r[0] for r in bad))


def truth_book(Wd):
    T = Wd["truth"]
    days = [d for d in sim.ALL_DAYS if d <= W.LAST_CLOSE]
    n = len(days)
    acc = {PID[c] or c: {"oh": T[c]["OH"][:n].astype(float), "ship": T[c]["S"][:n].astype(float),
                         "recv": T[c]["R"][:n].astype(float)} for c in T}
    i6, i31 = sim.DIDX[W.LAST_CLOSE], sim.DIDX[dt.date(2026, 12, 31)]
    notices = {PID[c]: float(T[c]["R"][i6 + 1:i31 + 1].sum()) for c in W.CURRENT}
    steps = [W.RUSHES[r][2] for r in ("R2019", "R2025A", "R2025B")]
    return calc.Book(days, acc, notices, steps, [PID[c] for c in W.CURRENT], W.LAST_CLOSE,
                     announced=list(W.ANNOUNCED.values()))


def calls(B):
    return {
        "fade": calc.rival_decay(B), "time_back": calc.rival_time_back(B),
        "book_burn": calc.rival_book_burn(B, "slope"), "book_burn_totals": calc.rival_book_burn(B, "total"),
        "monthly_fall": calc.rival_monthly_fall(B), "per_account_fall": calc.rival_per_account_monthly_fall(B),
        "per_account_fade": calc.rival_per_account_decay(B), "no_payback": calc.rival_no_payback(B),
        "scaled_nov": calc.rival_scaled(B, "nov"), "scaled_sepnov": calc.rival_scaled(B, "sepnov"),
        "scaled_dec": calc.rival_scaled(B, "dec"), "holt_winters": calc.rival_holt_winters(B),
        "trend": calc.rival_trend(B), "dwell": calc.rival_dwell(B)}


def backtest(B, parts, horizon=90):
    """Leave one earlier rush out: for each account with two or more episodes, predict its excess path
    after the left-out step from the other episodes by the burn (share of cumulative shipments), the
    fade (half-life) and the calendar rate (pallets a month); mean absolute error over the days scored."""
    tot = {"burn": 0.0, "fade": 0.0, "rate": 0.0}
    per = {}
    for code in B.current:
        p = parts[code]
        eps = {e[0]: e for e in p["episodes"]}
        if len(eps) < 2:
            continue
        at = calc.daily_normal(B, code, p["cover"], dict(p["S"]))
        a = B.acc[code]
        paths = {}
        for st in eps:
            i0 = B.idx[st]
            e0 = float(np.mean([a["oh"][i] - at(B.days[i]) for i in range(i0 - 7, i0)]))
            later = [d for d in B.announced if d > st]
            stop = B.idx.get(min(later)) if later else 10 ** 9
            es, cs, ts, cum = [], [], [], 0.0
            for i in range(i0, min(i0 + horizon, stop, len(B.days))):
                cum += a["ship"][i]
                es.append(a["oh"][i] - at(B.days[i]))
                cs.append(cum)
                ts.append(i - i0 + 1)
            paths[st] = (e0, np.array(es), np.array(cs), np.array(ts))
        for st in eps:
            others = [o for o in eps if o != st]
            w = np.array([eps[o][4] for o in others])
            s = float(np.dot(w, [eps[o][1] for o in others]) / w.sum())
            hls, rates = [], []
            for o in others:
                e0o, eso, _, tso = paths[o]
                ok = eso > 0.05 * e0o
                if ok.sum() >= 10:
                    b = np.polyfit(tso[ok] / 30.4, np.log(eso[ok]), 1)[0]
                    hls.append(np.log(0.5) / b if b < 0 else 50.0)
                rates.append(eps[o][1] * eps[o][4] / max(1, eps[o][3]) * 30.4)
            e0, es, cs, ts = paths[st]
            act = np.maximum(es, 0.0)
            preds = {"burn": np.maximum(0.0, e0 - s * cs),
                     "fade": e0 * 0.5 ** (ts / 30.4 / float(np.median(hls))) if hls else e0 + 0 * ts,
                     "rate": np.maximum(0.0, e0 - float(np.median(rates)) * ts / 30.4)}
            for k, pr in preds.items():
                err = float(np.abs(pr - act).mean())
                tot[k] += err
                per[(st, k)] = per.get((st, k), 0.0) + err
    return tot, per


def flip_scale(B, lo, hi, pred, it=12):
    """The share multiplier between lo and hi where pred(golden) changes, by bisection."""
    flo = pred(calc.golden(B, opt={"scale": lo}))
    assert flo != pred(calc.golden(B, opt={"scale": hi}))
    for _ in range(it):
        mid = 0.5 * (lo + hi)
        if pred(calc.golden(B, opt={"scale": mid})) == flo:
            lo = mid
        else:
            hi = mid
    return 0.5 * (lo + hi)


def a2_hits(inp, gT, gS, **kw):
    t, s = asks.a2_rows(inp["receipts"], inp["manifest"], inp["LA"], inp["LB"], inp["to_local"], **kw)
    th = sum(1 for k in asks.MONTHS_2026 if int(round(t[k])) == gT[k])
    sh = sum(1 for k in asks.MONTHS_2026 if abs(round(s[k], 3) - gS[k]) < 0.0005)
    return th, sh


def headers(target):
    import openpyxl
    out = {}
    for f in sorted(os.listdir(target)):
        p = os.path.join(target, f)
        if f.endswith(".csv"):
            with open(p, newline="") as fh:
                out[f] = next(csv.reader(fh))
        elif f.endswith(".txt"):
            with open(p, newline="") as fh:
                out[f] = fh.readline().strip().split("|")
        elif f.endswith(".xlsx"):
            wb = openpyxl.load_workbook(p, read_only=True)
            cells = []
            for ws in wb.worksheets:
                for r in ws.iter_rows(min_row=1, max_row=6, values_only=True):
                    cells += [str(x) for x in r if isinstance(x, str)]
            out[f] = cells
    return out


def doc_texts(target):
    out = {}
    for f in sorted(os.listdir(target)):
        if f.endswith(".pdf"):
            out[f] = re.sub(r"\s+", " ", read.pdf_text(os.path.join(target, f)))
    return out


def run(target, golden_dir, Bt=None):
    S = Sheet()
    Wd = Bt.Wd if Bt is not None and hasattr(Bt, "Wd") else __import__("core").build_world()
    B = read.book(target)
    g, parts = calc.golden(B, return_parts=True)
    TB = truth_book(Wd)
    tg = calc.golden(TB)
    R = calls(B)
    TR = {"fade": calc.rival_decay(TB), "book_burn": calc.rival_book_burn(TB, "slope"),
          "book_burn_totals": calc.rival_book_burn(TB, "total"),
          "per_account_fall": calc.rival_per_account_monthly_fall(TB)}
    use = g["use"]

    # ------------------------------------------------------------------ R: the three rules executed
    S.ok("R1a", B.days[0] == dt.date(2019, 4, 1) and B.last == W.LAST_CLOSE and len(B.days) == 2807 and
         all(B.acc[c]["oh"][B.idx[B.last]] > 0 for c in B.current) and len(B.current) == 14,
         "the golden reads each account's daily closing stock and shipments, 1 April 2019 to 6 December 2026, "
         "from the legacy, nightly and outbound exports (2,807 days, 14 accounts under contract)")
    banned = re.compile(r"excess|burn|normal stock|forecast|overflow|\bcover\b|half.?life|bought.?ahead|block", re.I)
    hdr = headers(target)
    hits = {f: [h for h in hs if banned.search(h)] for f, hs in hdr.items() if f not in PK.DISTRACTORS}
    S.ok("R1b", not any(hits.values()), "no shipped column or header holds an excess, a burn share, a normal stock, "
         "a forecast, an overflow or a block count: %s" % {f: v for f, v in hits.items() if v})
    S.ok("R1c", all(len(parts[c]["episodes"]) >= 2 for c in SEASONAL),
         "each of the five seasonal accounts has two or more earlier rushes in its own rows: %s" %
         {c: len(parts[c]["episodes"]) for c in SEASONAL})
    mis = [("R2.1", "the legacy months' mean over the rows on file (Sundays not carried)", {"sundays": "rows"}),
           ("R2.2", "the second Rosales legacy code left unmapped", {"legacy_codes": "first"}),
           ("R2.3", "the parallel run of 12 to 30 September 2022 stacked from both systems", {"stack_parallel": True}),
           ("R2.4", "the nightly run's UTC stamp date read as the close", {"shift_run_day": 1}),
           ("R2.6a", "box-only notices read as no pallets", {"notice_basis": "zero"}),
           ("R2.6b", "box-only notices at the book's pallets per box", {"notice_basis": "book"}),
           ("R2.6c", "box-only notices at the account's pallets per receipt row", {"notice_basis": "rows"})]
    for label, what, kw in mis:
        r = calc.golden(read.book(target, **kw))
        S.ok(label, r["blocks"] == 2, "%s: %d blocks (block uses %s)" % (what, r["blocks"],
                                                                         [round(u, 3) for u in r["use"][:3]]))
    r = calc.golden(B, opt={"population": "all"})
    S.ok("R2.5", r["blocks"] == 2 and all(B.acc[c]["oh"][B.idx[B.last]] == 0 for c in B.acc if c not in B.current),
         "the two former accounts kept in the population: %d blocks; they hold nothing at the close" % r["blocks"])
    S.ok("R3a", tg["blocks"] == 2 and calc.rival_decay(TB)["blocks"] == 0 and R["fade"]["blocks"] == 0,
         "the default path (the book's excess faded at the earlier rushes' half-life) on the clean truth files "
         "0 blocks (none) with every filed rule applied; on the pack also 0 (half-life %.2f months)" %
         R["fade"]["half_life_months"])
    lower = {k: TR[k]["blocks"] for k in ("book_burn", "book_burn_totals", "per_account_fall")}
    S.ok("R3b", tg["blocks"] == 2 and all(v != 2 for v in lower.values()),
         "on the clean truth the decisive step alone (each account's excess burned at its own share) files 2; "
         "each lower rung's single correction stays wrong: %s" % lower)
    plant = []
    for c in B.acc:
        oh = B.acc[c]["oh"]
        seen = {}
        for i in range(0, len(oh) - 14):
            win = tuple(oh[i:i + 14])
            if len(set(win)) <= 2:
                continue
            if win in seen and i - seen[win] >= 14:
                plant.append((c, i))
            seen.setdefault(win, i)
    S.ok("R3c", not plant, "no plant: no copied or repeated 14-day block of any account's daily stock, no memo, "
         "draft or stakeholder figure argues a number of blocks (%d repeats)" % len(plant))

    # ------------------------------------------------------------------ L: the ladder and the margins
    S.ok("L1", g["blocks"] == 2 and use[0] == 1.0 and 0.90 <= use[1] <= 1.0 and use[2] <= 0.25,
         "the golden on the pack: 2 blocks, block uses %s" % [round(u, 4) for u in use[:4]])
    S.ok("L2", tg["blocks"] == 2, "the golden on the clean truth: 2 blocks")
    want = {"fade": 0, "time_back": 0, "book_burn": 0, "book_burn_totals": 1, "monthly_fall": 0,
            "per_account_fall": 1, "per_account_fade": 0, "no_payback": 3, "scaled_nov": 0, "dwell": 0}
    got = {k: R[k]["blocks"] for k in R}
    S.ok("L3", all(got[k] == v for k, v in want.items()) and got["holt_winters"] >= 3 and got["trend"] >= 3,
         "every other route by name: %s" % got)
    S.ok("L4", got["scaled_dec"] == 2 and got["scaled_sepnov"] != 2,
         "the accidental landing measured: last year's January to June scaled by the projected close of 31 "
         "December over December 2025 files 2 (block 3 at %.2f); by September to November, %d" %
         (R["scaled_dec"]["use"][2], got["scaled_sepnov"]))
    up = flip_scale(B, 1.0, 2.0, lambda r: r["blocks"] >= 2)
    dn = flip_scale(B, 0.05, 1.0, lambda r: r["blocks"] <= 2)
    S.ok("L5", up >= 1.2 and dn <= 0.6, "flip: block 2 falls under three quarters only with every account's share "
         "%.0f per cent higher; block 3 reaches it only with every share %.0f per cent lower" %
         (100 * (up - 1), 100 * (1 - dn)))
    grid = []
    for sh, cv, fc in itertools.product(["weighted", "mean", "median", "latest", "pooled", ("weighted", 0.85),
                                         ("weighted", 1.15)], ["median", "mean_2024"], ["same_month", "two_year", "flat"]):
        opt = {"cover": cv, "forecast": fc}
        if isinstance(sh, tuple):
            opt.update(share=sh[0], scale=sh[1])
        else:
            opt["share"] = sh
        r = calc.golden(B, opt=opt)
        grid.append((r["blocks"], r["use"][1], r["use"][2]))
    S.ok("L6", all(x[0] == 2 for x in grid) and min(x[1] for x in grid) >= 0.80 and max(x[2] for x in grid) <= 0.35,
         "grid: %d cells (share basis x normal cover x shipment forecast) all file 2; block 2 %.2f to %.2f, block 3 "
         "%.2f to %.2f" % (len(grid), min(x[1] for x in grid), max(x[1] for x in grid), min(x[2] for x in grid),
                           max(x[2] for x in grid)))
    alone = {str(st): calc.golden(B, opt={"steps": [st]})["blocks"] for st in B.steps}
    S.ok("L7", all(v == 2 for v in alone.values()), "each earlier rush alone files 2: %s" % alone)
    tot, per = backtest(B, parts)
    apr = {k: per[(dt.date(2025, 4, 5), k)] for k in ("burn", "fade", "rate")}
    S.ok("L8", tot["burn"] * 1.1 <= min(tot["fade"], tot["rate"]) and apr["burn"] * 1.4 <= min(apr["fade"], apr["rate"]),
         "out of sample on the earlier rushes (each left out in turn, 90 days scored): mean absolute error summed over "
         "accounts, burn %.0f, fade %.0f, calendar rate %.0f; the April 2025 step alone %.0f, %.0f, %.0f" %
         (tot["burn"], tot["fade"], tot["rate"], apr["burn"], apr["fade"], apr["rate"]))
    jan = calc.WINDOW[0]
    E31 = calc._book_parts(B)["E"]
    lam = 0.5 ** (1.0 / R["fade"]["half_life_months"])
    fade_fall_jan = E31 * (1 - lam ** (31 / 30.4))
    seas_ship_jan = sum(parts[c]["fc"][jan] for c in SEASONAL)
    seas_e = sum(calc_excess(B, parts, c) for c in SEASONAL)
    S.ok("L9", seas_e > 0.8 * E31 and fade_fall_jan - (E31 - seas_e) > seas_ship_jan,
         "the fade needs bought-ahead stock to leave faster than its owners ship: it takes %.0f pallets of excess out "
         "in January, more than the %.0f the accounts outside the five hold plus the five's own January shipments of "
         "%.0f; the five hold %.0f of the %.0f pallets of excess at 31 December" %
         (fade_fall_jan, E31 - seas_e, seas_ship_jan, seas_e, E31))
    tot_e = sum(calc_excess(B, parts, c) for c in B.current)
    S.ok("L10", abs(tot_e - E31) / E31 < 0.03, "lens: the default path and the golden start from one population at one "
         "moment (the book's excess at 31 December %.0f against the accounts' summed %.0f); they differ only in how it "
         "leaves" % (E31, tot_e))

    # ------------------------------------------------------------------ SF: the semantic fork table
    txt = doc_texts(target)
    pol, off, adv, dic = txt[PK.F_POLICY], txt[PK.F_OFFER], txt[PK.F_ADVISORIES], txt[PK.F_DICT]
    S.ok("SF1", "Blocks are let whole and for the whole term" in off and "whole, and for the whole term" in pol,
         "a block is taken for the whole term: the offer and policy 4.1")
    r = calc.golden(B, opt={"rack": 39740})
    S.ok("SF2", r["blocks"] == 2, "overflow above the racks' practical fill (39,740): %d blocks" % r["blocks"])
    S.ok("SF3", calc.golden(B, opt={"population": "all"})["blocks"] == 2 and
         not any(B.notices.get(c) for c in B.acc if c not in B.current), "every account in the history: 2 blocks")
    r1, r2 = calc.golden(B, opt={"reading": "month_end"}), calc.golden(B, opt={"reading": "peak"})
    S.ok("SF4", r1["blocks"] == 2 and r2["blocks"] == 2 and "mean of each day's closing stock in the month" in pol,
         "a month's need at the month-end close %d, at the peak day %d; policy 3.1 files the mean" %
         (r1["blocks"], r2["blocks"]))
    r = calc.golden(B, opt={"test": "set"})
    S.ok("SF5", r["blocks"] == 2, "the block test read as the blocks' average use together: %d" % r["blocks"])
    S.ok("SF6", "in our racks or outside them" in pol, "racked pallets only would file 0; policy 3.1 counts every "
         "pallet in our racks or outside them")
    r = calc.golden(B, opt={"close_shift": -1})
    S.ok("SF7", r["blocks"] == 2 and "holds the balances at the close of the business day just ended" in dic,
         "the close read from the run stamped 6 December (the close of 5 December): %d; the dictionary's run line" %
         r["blocks"])
    r = calc.golden(B, opt={"dec_rest": "last_year"})
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(target, PK.F_TRACKER), read_only=True)
    trk = " ".join(str(c) for row in wb["Dec arrivals"].iter_rows(min_row=1, max_row=3, values_only=True)
                   for c in row if c)
    S.ok("SF8", r["blocks"] == 1 and "arrivals due after the close of Sunday 6 December" in trk,
         "the rest of December at last year's net flow files %d; the tracker's header files the notices as the "
         "arrivals due after the close" % r["blocks"])
    S.ok("SF9", calc.golden(read.book(target, notice_basis="book"))["blocks"] == 2 and
         "that account's own pallets per box received over the last 90 days" in trk,
         "box-only notices at the book's pallets per box: 2; the tracker's planning basis files the account's own")
    r = calc.golden(B, opt={"steps": [dt.date(2019, 8, 16), dt.date(2025, 4, 3), dt.date(2025, 8, 1)]})
    S.ok("SF10", r["blocks"] == 2 and B.future_steps == [dt.date(2027, 1, 1)] and len(B.steps) == 3,
         "the earlier steps dated by the advisories that announced them: %d; the step itself is 1 January 2027" %
         r["blocks"])
    r = calc.golden(B, opt={"steps": list(B.steps) + [dt.date(2022, 10, 1)]})
    S.ok("SF11", r["blocks"] == 2, "a surge with no duty step (the 2022 congestion unwind) taken as a fourth rush: %d" %
         r["blocks"])
    r = calc.golden(B, opt={"test": "days"})
    S.ok("SF12", r["blocks"] == 2, "a block's position-months counted in position-days: %d" % r["blocks"])

    # ------------------------------------------------------------------ A: the asks
    inp = read.ask_inputs(target)
    a1 = asks.a1_golden(inp["site_onhand"], inp["lpn"])
    lpn_only = asks.a1_lpn_only(inp["lpn"])
    prev = read.site_onhand(target, inp["close"] - dt.timedelta(days=1))
    S.ok("A1a", a1 == {"CAR1": 1815, "RDZ1": 1589, "WIL1": 1042, "CMP1": 776},
         "pallets outside the racks at the close of 6 December: %s" % a1)
    S.ok("A1b", all(lpn_only[s] != a1[s] for s in a1) and prev == inp["site_onhand"] and inp["close"].weekday() == 6,
         "the pallet export's non-rack rows %s miss every building; the close of Sunday 6 December equals Saturday's, "
         "so the run read for the close does not move A1" % lpn_only)
    split = Wd["split_asof"]
    trl = {s: sum(p["TRL:" + s] for p in split.values()) for s in asks.BUILDINGS}
    S.ok("A1c", all(a1[s] - lpn_only[s] == trl[s] for s in a1),
         "the gap is the yard's storage trailers, which only the nightly site balance counts: %s" % trl)
    teu, share = asks.a2_golden(inp["receipts"], inp["manifest"], inp["LA"], inp["LB"], inp["to_local"])
    gT, gS = asks.rounded_a2(teu, share)
    e_t = min(0.5 - abs(teu[k] - round(teu[k])) for k in teu)
    e_s = min(0.0005 - abs(share[k] - round(share[k], 3)) for k in share)
    S.ok("A2a", e_t >= 0.25 and e_s >= 0.00025, "bin edges: TEU %.2f from the nearest whole-TEU edge, shares %.5f "
         "from the nearest edge at three decimals (a quarter cell or more)" % (e_t, e_s))
    stops = {"rows, no trailers": dict(dedup="none", trailers="exclude"),
             "box numbers collapsed, no trailers": dict(dedup="global", trailers="exclude"),
             "box numbers collapsed, manifest": dict(dedup="global", trailers="manifest"),
             "rows, manifest": dict(dedup="none", trailers="manifest"),
             "visits, no trailers": dict(dedup="visit", trailers="exclude"),
             "visits, trailers at 53/20 TEU": dict(dedup="visit", trailers="len20"),
             "visits, manifest, UTC months": dict(dedup="visit", trailers="manifest", month="utc"),
             "visits, manifest, 45ft at 2 TEU": dict(dedup="visit", trailers="manifest", teu45=2.0)}
    res = {k: a2_hits(inp, gT, gS, **v) for k, v in stops.items()}
    gold = a2_hits(inp, gT, gS, dedup="visit", trailers="manifest")
    S.ok("A2b", gold == (10, 10), "the golden path reproduces all twenty figures: %s" % (gold,))
    S.ok("A2c", all(v[0] <= 1 and v[1] <= 1 for k, v in res.items() if "UTC" not in k and "45ft" not in k) and
         res["visits, manifest, UTC months"][0] <= 6 and res["visits, manifest, 45ft at 2 TEU"][0] == 0,
         "every stop misses (TEU hits, share hits of ten): %s" % res)
    cr, _, _, _ = cracker.score(inp)
    p1 = (cr["C1"]["score"] + cr["M"]["score"]) / 2
    p2 = (cr["C2"]["score"] + cr["M"]["score"]) / 2
    S.ok("A3", max(p1, p2) <= 40.0 and cr["M"]["kept"] < 0.35 and cr["C1"]["a1"] == 0 and cr["C1"]["teu"] == 0,
         "cracker sheet on the pack: pair ceiling C1 with the mirror %.1f, C2 with the mirror %.1f; mirror leakage "
         "%.1f per cent" % (p1, p2, 100 * cr["M"]["kept"]))
    pairs = {}
    for r in Wd["receipts"]:
        if r.get("part"):
            pairs.setdefault((r["equipment"], r["account"]), []).append(r["day"])
    span = [ds for ds in pairs.values() if len(ds) == 2]
    same_month = all((a.year, a.month) == (b.year, b.month) or (b - a).days <= 3 for a, b in span)
    cross = sum(1 for a, b in span if (a.year, a.month) != (b.year, b.month))
    S.ok("A4", len(span) > 1000 and same_month,
         "separation: %d boxes devanned over two shifts, each pair within three days (%d across a month end, where "
         "the call's month means move by the second part's pallets for a day); the call reads only each account's "
         "daily totals" % (len(span), cross))
    tl = {m["trailer"] for m in inp["manifest"]}
    rc_tl = {r["equipment"] for r in inp["receipts"] if r["type"] == "53DV" and r["equipment"].startswith("AVTL")}
    S.ok("A5", tl == rc_tl, "every Avila trailer in the receipts is in the partner's manifest and back (%d)" % len(tl))

    # ------------------------------------------------------------------ G: the input gates
    meta = json.load(open(os.path.join(os.path.dirname(target), "metadata.json")))
    files = sorted(os.listdir(target))
    fmts = sorted({f.rsplit(".", 1)[1] for f in files})
    big = max(x["rows"] or 0 for x in meta["files"])
    S.ok("G1", len(files) >= 10 and len(fmts) >= 3 and big >= 25000,
         "%d files, %d formats %s, largest %d rows" % (len(files), len(fmts), fmts, big))
    S.ok("G2", len(meta["distractor_files"]) >= 2 and all(d in files for d in meta["distractor_files"]) and
         not any("distractor" in f.lower() for f in files), "distractors named in metadata.json only: %s" %
         meta["distractor_files"])
    for k, d in enumerate(meta["distractor_files"]):
        tmp = tempfile.mkdtemp(dir=os.path.dirname(target), suffix=".tmp")
        try:
            for f in files:
                if f != d:
                    os.symlink(os.path.join(target, f), os.path.join(tmp, f))
            r = calc.golden(read.book(tmp))
            ii = read.ask_inputs(tmp)
            t2, s2 = asks.rounded_a2(*asks.a2_golden(ii["receipts"], ii["manifest"], ii["LA"], ii["LB"],
                                                     ii["to_local"]))
            S.ok("G3.%d" % (k + 1), r["blocks"] == 2 and r["use"] == g["use"] and t2 == gT and s2 == gS and
                 asks.a1_golden(ii["site_onhand"], ii["lpn"]) == a1, "%s deleted: the call and every ask unchanged" % d)
        finally:
            shutil.rmtree(tmp)

    # ------------------------------------------------------------------ S: the realism sweep
    ms = {}
    for c in B.acc:
        P_, S_, _ = B.monthly(c)
        ms[c] = tuple(round(v) for k, v in sorted(S_.items()) if k >= (2023, 1))
    S.ok("S1", len(set(ms.values())) == len(ms), "no two accounts share a monthly shipment series")
    bk, _, _ = calc.book_series(B, list(B.acc))
    vals = [round(v) for v in bk.values()]
    S.ok("S2", len(set(vals)) == len(vals), "no two months of the book's stock are equal")
    figs = list(a1.values()) + list(gT.values())
    S.ok("S3", not any(v % 100 == 0 for v in figs) and len(set(gS.values())) == 10,
         "no graded count is a round hundred; the ten shares are distinct")
    dist = {c: len(set(B.acc[c]["ship"][B.idx[dt.date(2025, 1, 1)]:])) for c in B.current}
    S.ok("S4", min(dist.values()) > 40, "no constant stream: each account ships at least %d distinct daily counts" %
         min(dist.values()))
    ok_cd = True
    for r in inp["receipts"]:
        if r["type"] in asks.TEU:
            cn = r["equipment"]
            ok_cd &= len(cn) == 11 and records.iso6346_check(cn[:10]) == int(cn[10])
    for m in inp["manifest"]:
        ok_cd &= records.iso6346_check(m["container"][:10]) == int(m["container"][10])
    ok_ss = all(records.sscc(int(l["lpn"][8:17]), int(l["lpn"][0:2]))[-1] == l["lpn"][-1] for l in inp["lpn"][:5000])
    S.ok("S5", ok_cd and ok_ss, "every container number carries its ISO 6346 check digit; every SSCC its GS1 check digit")
    people = {"Mariah Norman", "Jerome Lam", "Nicole Freeman", "Brandon Page"}
    named = set(re.findall(r"\b(Mariah Norman|Jerome Lam|Nicole Freeman|Brandon Page)\b", " ".join(txt.values())))
    creators = set()
    for f in files:
        if f.endswith(".xlsx"):
            with zipfile.ZipFile(os.path.join(target, f)) as z:
                core = z.read("docProps/core.xml").decode()
                creators |= set(re.findall(r"<dc:creator>([^<]*)</dc:creator>", core))
                creators |= set(re.findall(r"<cp:lastModifiedBy>([^<]*)</cp:lastModifiedBy>", core))
                app = z.read("docProps/app.xml").decode()
                S.ok("S6.%s" % f[:14], "openpyxl" not in core + app and "Microsoft Excel" in app,
                     "%s carries Excel's own properties, no writer name" % f)
    roles = {"Capacity Planning", "Avila Dispatch"}
    S.ok("S7", named <= people and creators <= people | roles,
         "personas from guard.py names only: documents %s, workbook authors %s" % (sorted(named), sorted(creators)))
    pdf_dates = {}
    for f in files:
        if f.endswith(".pdf"):
            b = open(os.path.join(target, f), "rb").read()
            pdf_dates[f] = re.findall(rb"D:(\d{8})", b)
            S.ok("S8.%s" % f[:12], b"ReportLab" not in b and all(x <= b"20261207" for x in pdf_dates[f]),
                 "%s: no writer banner, dated %s" % (f, sorted({x.decode() for x in pdf_dates[f]})))
    mod = []
    for f in files:
        if f.endswith(".xlsx"):
            with zipfile.ZipFile(os.path.join(target, f)) as z:
                mod += re.findall(r"<dcterms:modified[^>]*>([^<]*)<", z.read("docProps/core.xml").decode())
    S.ok("S9", all(m <= "2026-12-07T23:59:59Z" for m in mod) and
         len({int(os.path.getmtime(os.path.join(target, f))) for f in files}) == 1,
         "no document dated after the as-of; one export time on every file")
    rd = {}
    for r in inp["receipts"]:
        d = read.pacific_date(r["received_at"])
        rd[d] = rd.get(d, 0) + 1
    ob = {}
    with open(os.path.join(target, PK.F_OUTBOUND), newline="") as fh:
        for r in csv.DictReader(fh):
            if r["ship_date"] >= "2025-01-01":
                ob[r["ship_date"]] = ob.get(r["ship_date"], 0) + int(r["pallets_shipped"])
    wk = [v for d, v in ob.items() if dt.date.fromisoformat(d).weekday() < 5]
    sa = [v for d, v in ob.items() if dt.date.fromisoformat(d).weekday() == 5]
    S.ok("S11", round(np.mean(sa) / np.mean(wk), 2) == calc.SAT, "the projection's Saturday weight %.2f is the book's own "
         "Saturday shipments over a weekday's in 2025 and 2026 (%.3f)" % (calc.SAT, np.mean(sa) / np.mean(wk)))
    S.ok("S10", len(set(rd.values())) > 60 and max(rd.values()) < 3 * np.median(list(rd.values())),
         "receipts per day vary (%d distinct daily counts)" % len(set(rd.values())))

    man = []
    with open(os.path.join(target, PK.F_MANIFEST), newline="") as fh:
        man = [r["file"] for r in csv.DictReader(fh)]
    S.ok("S12", sorted(man) == sorted(f for f in files if f != PK.F_MANIFEST),
         "the export manifest lists every shipped file but itself (%d)" % len(man))
    late = set()
    iso = re.compile(rb"20\d\d-\d\d-\d\d")
    for f in files:
        if f.endswith((".csv", ".txt")):
            if any(x.decode() > "2026-12-07" for x in iso.findall(open(os.path.join(target, f), "rb").read())):
                late.add(f)
    import openpyxl as _ox
    for f in files:
        if f.endswith(".xlsx"):
            wb_ = _ox.load_workbook(os.path.join(target, f), read_only=True)
            for ws_ in wb_.worksheets:
                for row in ws_.iter_rows(values_only=True):
                    if any(isinstance(v, dt.datetime) and v.date() > AS_OF for v in row):
                        late.add(f)
                        break
    allow = {PK.F_TRACKER, PK.F_MANIFEST, PK.F_SITES}
    S.ok("S13", late <= allow, "dates after the as-of sit only in forward-looking records (the December arrivals, "
         "the sublease's end): %s" % sorted(late))
    gstamps = []
    for f in sorted(os.listdir(golden_dir)):
        b = open(os.path.join(golden_dir, f), "rb").read()
        if f.endswith(".pdf"):
            gstamps += [x.decode() for x in re.findall(rb"D:(\d{8})", b)]
        else:
            with zipfile.ZipFile(os.path.join(golden_dir, f)) as z:
                core = z.read("docProps/core.xml").decode()
                gstamps += [x[:10].replace("-", "") for x in re.findall(r"<dcterms:\w+[^>]*>([^<]*)<", core)]
                S.ok("S14", "openpyxl" not in core + z.read("docProps/app.xml").decode(),
                     "the golden workbook carries Excel's properties")
    S.ok("S15", gstamps and all(x == "20261207" for x in gstamps), "the goldens are stamped 7 December 2026: %s" %
         sorted(set(gstamps)))

    # ------------------------------------------------------------------ P: pins once, counter-pins none
    pins = {"in our racks or outside them": PK.F_POLICY, "at least three quarters of its position-months": PK.F_POLICY,
            "holds the balances at the close of the business day just ended": PK.F_DICT,
            "L5G1 2.25": PK.F_DICT, "Blocks are let whole and for the whole term": PK.F_OFFER}
    where = {p: [f for f, t in txt.items() if p in t] for p in pins}
    S.ok("P1", all(where[p] == [f] for p, f in pins.items()),
         "each load-bearing sentence sits in exactly one file: %s" % where)
    counter = re.compile(r"half.?life|fade|burn|weeks of supply|net(ted|ting)? against|bought.?ahead|per account|"
                         r"forecast method|regress|decay", re.I)
    found = {f: counter.findall(t) for f, t in txt.items()}
    S.ok("P2", not any(found.values()), "no shipped sentence names a method for the stock: %s" %
         {f: v for f, v in found.items() if v})
    S.ok("P3", not re.search(r"\b(?:recommend|should take|we expect to need)\b", " ".join(txt.values()), re.I),
         "no document recommends or forecasts a number of blocks")

    # ------------------------------------------------------------------ F: the goldens and figures.json
    F = json.load(open(os.path.join(HERE, "figures.json")))
    import openpyxl as ox
    gw = ox.load_workbook(os.path.join(golden_dir, "overflow_blocks_h1_2027.xlsx"), data_only=True)
    a1x = {r[1]: r[4] for r in gw["Off-rack 6 Dec"].iter_rows(min_row=5, values_only=True) if r[1]}
    a2x = {r[0]: (r[1], r[4]) for r in gw["Port TEU 2026"].iter_rows(min_row=5, values_only=True) if r[0]}
    mon = {k: "%s 2026" % sim_month(k) for k in asks.MONTHS_2026}
    S.ok("F1", a1x == F["a1"] == a1, "the workbook's off-rack sheet, figures.json and the recompute agree")
    S.ok("F2", all(a2x[mon[k][:3] + " 2026"] == (gT[k], round(gS[k], 3)) for k in asks.MONTHS_2026) and
         all(F["a2_teu"][mon[k]] == gT[k] and F["a2_share"][mon[k]] == "%.3f" % gS[k] for k in asks.MONTHS_2026),
         "the workbook's port sheet, figures.json and the recompute agree on all twenty figures")
    call_cell = gw["Blocks"]["A4"].value
    S.ok("F3", F["call"]["blocks"] == 2 and "take 2 of the 6 blocks" in call_cell and
         [round(u, 4) for u in g["use"]] == [round(x[4], 4) for x in gw["Blocks"].iter_rows(min_row=16, max_row=21,
                                                                                               values_only=True)],
         "the workbook's call line and block table: %s" % call_cell)
    S.report()
    return S


def sim_month(k):
    return ["", "January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
            "November", "December"][k[1]]


def calc_excess(B, parts, c):
    p = parts[c]
    S_ext = dict(p["S"])
    S_ext.update({k: v for k, v in p["fc"].items() if v is not None})
    at = calc.daily_normal(B, c, p["cover"], S_ext)
    return max(0.0, p["oh_dec31"] - at(dt.date(2026, 12, 31)))


if __name__ == "__main__":
    task = os.path.dirname(HERE)
    run(os.path.join(task, "target"), os.path.join(task, "golden"))
