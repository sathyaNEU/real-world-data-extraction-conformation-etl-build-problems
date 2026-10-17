"""The generator's assertions for task469: the three rules executed on the clean truth (R1 to R3c), every row of the
semantic fork table (SF1 to SF12), the ladder and its margins, the Layer 1 items' neutrality, the schedule ask's device
and stops, the separation of the call from the ask, the distractors, the golden against figures.json, the input gates
and the realism sweep.

    python3 task469/generator/checks.py      (after build.py has written target/ and golden/)
"""
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import numpy as np                # noqa: E402
import pandas as pd               # noqa: E402

import bundle as B                # noqa: E402
import golden as GO               # noqa: E402
import methods as M               # noqa: E402
import pack as P                  # noqa: E402
import schedule as SC             # noqa: E402
import ssim as SS                 # noqa: E402
import world as W                 # noqa: E402

TASK = os.path.dirname(HERE)
AS_OF = "2026-10-30"
DASH = re.compile("[–—]")
GOLD_BACK = {"S1", "S4", "S5"}
GOLD_SHORTEST = {"S1": 40, "S2": 40, "S3": 30, "S4": 45, "S5": 35, "S6": 30}


# ---------------------------------------------------------------- helpers
def verdicts(pop, filings, f=M.km):
    rv = M.review(pop, filings, f)
    return {s for s in W.ORDER if rv[s]["send_back"]}, {s: rv[s]["shortest"] for s in W.ORDER}, rv


def pdf_text(path):
    out = subprocess.run(["pdftotext", "-raw", path, "-"], capture_output=True, text=True, check=True).stdout
    return re.sub(r"\s+", " ", out)


def xlsx_text(path):
    with zipfile.ZipFile(path) as z:
        return " ".join(z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.endswith(".xml"))


def copy_without(target, keys):
    tmp = tempfile.mkdtemp()
    for f in os.listdir(target):
        if f not in {W.FILES[k] for k in keys}:
            shutil.copy(os.path.join(target, f), tmp)
    return tmp


def clean_truth(x):
    """The generator's own turns, before any record was written: the need of every turn known exactly, the stand from
    the allocation, the time each had in unrounded minutes."""
    stand = {(r["day"], r["reg"], round(r["actual"], 6)): r["stand"] for r in x["rows"] if r["kind"] == "ARR"}
    T = pd.DataFrame(x["turns"])
    T["stand"] = [stand[(d, rg, round(a, 6))] for d, rg, a in zip(T["day"], T["reg"], T["aibt"])]
    T = T[T["stand"].str.startswith("E") & T["day"].map(W.peak)].copy()
    T["exact"] = T["turn_min"] > 0
    T["need"] = np.where(T["exact"], T["B"] + T["turn_min"], np.nan)
    return T


def run(target, golden, x=None, meta_path=None, fig_path=None):
    bad = []

    def check(label, cond, detail=""):
        print("%s  %s%s" % ("PASS" if cond else "FAIL", label, ("  [" + str(detail) + "]") if detail != "" else ""))
        if not cond:
            bad.append(label)

    x = x or P.everything()
    F = json.load(open(fig_path or os.path.join(HERE, "figures.json"), encoding="utf-8"))
    meta = json.load(open(meta_path or os.path.join(TASK, "metadata.json"), encoding="utf-8"))
    files = sorted(os.listdir(target))
    reg = B.registers(target)
    fil = reg["filings"]
    T = B.turns_2026(target, reg)
    pop = B.population(T)
    back, shortest, rv = verdicts(pop, fil)
    fx = GO.compute(target)

    # ------------------------------------------------------------------ the call
    check("call: the product-limit estimate sends back %s and accepts the rest" % sorted(GOLD_BACK), back == GOLD_BACK,
          sorted(back))
    check("call: shortest acceptable ground times %s" % GOLD_SHORTEST, shortest == GOLD_SHORTEST, shortest)
    check("call: the filings are the ones the world files", fil == {s: W.SEGMENTS[s]["filed"] for s in W.ORDER}, fil)

    # ------------------------------------------------------------------ R1 to R3: the three rules on the clean truth
    C = clean_truth(x)
    true_share = {s: float((C[C["seg"] == s]["N"] > fil[s]).mean()) for s in W.ORDER}
    true_short = {s: next(g for g in W.GRID if (C[C["seg"] == s]["N"] > g).mean() <= W.STANDARD) for s in W.ORDER}
    check("R1 truth: every turn's need known, the share needing longer than the filing sends back %s; shortest %s" % (
        sorted(GOLD_BACK), GOLD_SHORTEST),
        {s for s in W.ORDER if true_share[s] > W.STANDARD} == GOLD_BACK and true_short == GOLD_SHORTEST,
        {s: round(v, 3) for s, v in true_share.items()})
    r0c = {s: M.r0_visible(C[C["seg"] == s], fil[s]) for s in W.ORDER}
    check("R2 analyst test on the clean truth: the overruns the records show over every turn accept all six (no share "
          "above %.3f): the default path's Ends on" % max(r0c.values()), max(r0c.values()) < W.STANDARD - 0.08,
          {s: round(v, 3) for s, v in r0c.items()})
    cc = {s: M.cond(C[C["seg"] == s], fil[s]) for s in W.ORDER}
    kc = {s: M.km(C[C["seg"] == s], fil[s]) for s in W.ORDER}
    check("R3a decisive step alone on the clean truth: the same records read as censored (fully observed turns, or the "
          "product-limit estimator) send back %s" % sorted(GOLD_BACK),
          {s for s in W.ORDER if cc[s] > W.STANDARD} == GOLD_BACK and {s for s in W.ORDER if kc[s] > W.STANDARD} == GOLD_BACK,
          {s: (round(cc[s], 3), round(kc[s], 3)) for s in W.ORDER})
    r1c = {s: float((C[C["seg"] == s]["aobt"] - C[C["seg"] == s]["aibt"] - C[C["seg"] == s]["hold_min"] >
                     fil[s]).mean()) for s in W.ORDER}
    check("R3b executing every filed rule with the stand time as the need sends back all six on the clean truth",
          min(r1c.values()) > 0.5, {s: round(v, 3) for s, v in r1c.items()})
    docs = {f: pdf_text(os.path.join(target, f)) for f in files if f.endswith(".pdf")}
    every = " ".join(docs.values()).lower()
    check("R3c no shipped document names censoring, a product-limit or survival estimate, a candidate verdict or a "
          "shortest ground time", not re.search(r"censor|product-limit|kaplan|survival", every))

    # ------------------------------------------------------------------ convergence of the estimate (C1)
    variants = {"E gates, peak (golden)": pop, "remote-stand turns added (SF2)": T[T["peak"]],
                "towed turns out (SF1)": pop[~pop["towed"]], "through stops out (SF3)": pop[~pop["through"]],
                "16 June to 14 September (SF7)": pop[pop["day"].between(dt.date(2026, 6, 16), dt.date(2026, 9, 14))]}
    lo_margin = []
    for vn, pv in variants.items():
        for en, f in M.ESTIMATORS.items():
            b2, s2, r2 = verdicts(pv, fil, f)
            ok = b2 == GOLD_BACK and s2 == GOLD_SHORTEST
            if not ok:
                check("C1 %s, %s: the same call" % (vn, en), False, (sorted(b2), s2))
            for s in W.ORDER:
                t = pv[pv["seg"] == s]
                g = GOLD_SHORTEST[s]
                lo_margin.append((abs(f(t, fil[s]) - W.STANDARD), s, vn, en, "filed"))
                lo_margin.append((f(t, g - 5) - W.STANDARD, s, vn, en, "g*-5"))
                lo_margin.append((W.STANDARD - f(t, g), s, vn, en, "g*"))
    worst = min(lo_margin)
    check("C1 SF1 SF2 SF3 SF7 SF12: five populations by four estimators (fully observed with < and <=, product-limit, "
          "censored lognormal) give the same six decisions and shortest times; nearest edge %.3f (%s %s %s at %s)" %
          worst, worst[0] >= 0.05)

    # stratified by the on-block hour (the time a turn had depends on the hour; the need does not)
    def strat(t, g, f=M.km):
        parts = [t[t["on"] < 13 * 60], t[t["on"] >= 13 * 60]]
        return sum(len(p) * f(p, g) for p in parts if len(p)) / len(t)
    bs = {s for s in W.ORDER if strat(pop[pop["seg"] == s], fil[s]) > W.STANDARD}
    ss = {s: M.gstar(pop[pop["seg"] == s], strat) for s in W.ORDER}
    check("C1 SF12 the product-limit estimate within on-block hour bands (before and after 13:00), weighted by turns: "
          "the same call and shortest times", (bs, ss) == (GOLD_BACK, GOLD_SHORTEST), (sorted(bs), ss))

    # ------------------------------------------------------------------ the semantic fork rows
    bge = {s for s in W.ORDER if M.cond(pop[pop["seg"] == s], fil[s]) > W.STANDARD}
    check("SF5 needing longer read as at least the ground time: the same call", bge == GOLD_BACK)
    near = min(abs(v - W.STANDARD) for vals in fx["curves"].values() for v in vals)
    check("SF6 no share within half a point of 20 per cent at any five-minute step (no more than against less than)",
          near > 0.005, round(near, 4))
    mism = int((T["arr_flight"].map(lambda n: B.block_of(reg["blocks"], n, dt.date(2026, 8, 1))) !=
                T["dep_flight"].map(lambda n: B.block_of(reg["blocks"], n, dt.date(2026, 8, 1)))).sum())
    check("SF8 every turn's arriving and departing flights belong to one relationship (departing flight's by SFG-3 2.1)",
          mism == 0, mism)
    # code 93 read as held: the turns whose not-ready minutes were all late-inbound minutes look ready in time
    T93 = B.turns_2026(target, dict(reg, groups=dict(reg["groups"], **{"93": "held for another reason"})))
    b93, s93, r93 = verdicts(B.population(T93), fil)
    check("SF9 code 93 read as held for another reason moves the call (to %s); ruled out by kdy_delay_code_groups.csv "
          "(93 aircraft not ready)" % sorted(b93), (b93, s93) != (GOLD_BACK, GOLD_SHORTEST), sorted(b93))
    Tns = turns_unsubcoded(target, reg)
    bns, sns, _ = verdicts(B.population(Tns), fil)
    check("SF10 Sable Coast's sub-coded delays left unread moves the call (to %s); ruled out by SFG-3 2.5" % sorted(bns),
          (bns, sns) != (GOLD_BACK, GOLD_SHORTEST), sorted(bns))
    b5 = {s for s in W.ORDER if M.r5_whole_delay(pop[pop["seg"] == s], fil[s]) > W.STANDARD}
    check("SF11 held minutes counted as need (R5) sends back %s; ruled out by SFG-3 2.4" % sorted(b5),
          b5 != GOLD_BACK and GOLD_BACK <= b5, sorted(b5))
    Top = turns_operator_clock(target, reg)
    bop, sop, _ = verdicts(B.population(Top), fil)
    check("SF4 the operators' own block times in place of the stand log's: the same call", (bop, sop) ==
          (GOLD_BACK, GOLD_SHORTEST), (sorted(bop), sop))

    # ------------------------------------------------------------------ the ladder
    rungs = M.rung_verdicts(pop, fil)
    rb = {k: {s for s in W.ORDER if v[s]["send_back"]} for k, v in rungs.items()}
    check("ladder R0 visible overruns over every turn accepts all six", rb["R0 visible overruns"] == set(),
          {s: round(v["share"], 3) for s, v in rungs["R0 visible overruns"].items()})
    check("ladder R1 stand time as need sends back all six", rb["R1 ground time as need"] == set(W.ORDER))
    check("ladder R2 like-for-like 2026 turns sends back Tellura/Pellmont and Grayling only",
          rb["R2 like-for-like 2026 turns"] == {"S4", "S6"}, sorted(rb["R2 like-for-like 2026 turns"]))
    check("ladder R4 one operator's type pooled across brands sends back Norvale/Pellmont only",
          rb["R4 operator pooled"] == {"S1"}, sorted(rb["R4 operator pooled"]))
    check("ladder R5 whole delay as need sends back four", rb["R5 whole delay as need"] == {"S1", "S2", "S4", "S5"},
          sorted(rb["R5 whole delay as need"]))
    sets = [frozenset(v) for v in rb.values()] + [frozenset(GOLD_BACK)]
    check("ladder every rung names its own set of send-backs", len(set(sets)) == len(sets))
    r0max = max(v["share"] for v in rungs["R0 visible overruns"].values())
    check("ladder R0 margin: largest visible-overrun share %.3f, at least 8 points under 20 per cent" % r0max,
          r0max <= W.STANDARD - 0.08)

    # ------------------------------------------------------------------ Layer 1 neutral
    uh = turns_hyphen_kept(target, reg)
    bh, shh, _ = verdicts(B.population(uh), fil)
    check("Layer 1 registrations read with and without the hyphen as written: towed turns split, the same call",
          (bh, shh) == (GOLD_BACK, GOLD_SHORTEST), int(len(pop) - len(B.population(uh))))
    check("Layer 1 the 2026 season schedule carries no DEI 050 leg and every leg's flight is in an operator's log",
          ssim26_clean(target), "")
    nb = turns_no_block_dates(target, reg)
    check("Layer 1 the 15 July block transfer (3450 to 3499) carries no Kelderby flight: the call unchanged",
          verdicts(B.population(nb), fil)[:2] == (GOLD_BACK, GOLD_SHORTEST))

    # ------------------------------------------------------------------ the schedule ask
    truth = SC.ask_a_truth(x["lines27"])
    got = B.ask_a(target, reg)
    check("ask: the pack's operating turns in the week of 12 July 2027 equal the operating lines' truth", got == truth,
          got)
    stops = {"every leg counted (the cracker)": dict(dedup="none"), "schedule fields deduplicated": dict(dedup="schedule"),
             "layover field as the day test": dict(night="layover"), "through stops missed": dict(through=False),
             "UTC files read as local": dict(utc_as_local=True)}
    moved = {}
    for k, h in stops.items():
        r = B.ask_a(target, reg, **h)
        moved[k] = sum(r[s][j] != truth[s][j] for s in W.ORDER for j in (0, 1))
    check("ask: counting every leg misses all 12 figures", moved["every leg counted (the cracker)"] == 12, moved)
    check("ask: deduplicating on schedule fields misses at least 10 of 12", moved["schedule fields deduplicated"] >= 10)
    check("ask: missing the through stops misses all 12", moved["through stops missed"] == 12)
    check("ask: reading the rotation layover (UTC days in Tellura's and Grayling's files) as the day test moves a "
          "figure", moved["layover field as the day test"] >= 1, moved["layover field as the day test"])
    t27 = SC.turns27(x["lines27"])
    short = [t for t in t27 if t["gt"] < W.SEGMENTS[t["seg"]]["filed"]]
    check("ask: the week's turns shorter than the filing are exactly the through stops (the golden's note, back-tested)",
          all(t["through"] for t in short) and all(t["gt"] < W.SEGMENTS[t["seg"]]["filed"] for t in t27 if t["through"]),
          len(short))
    legs, refs = B.legs_2027(target)
    dup_keys = set(refs)
    sched_twins = twins_in(legs, refs)
    check("ask: genuine same-schedule pairs (two aircraft) exist beside the duplicates, so a schedule dedup deletes "
          "real turns", sched_twins >= 2, sched_twins)
    check("ask: every relationship carries legs a partner sells (the device moves every row)",
          all(any(r for r in refs if B.block_of(reg["blocks"], r[1], dt.date(2027, 7, 12))[1] == W.SEGMENTS[s]["op"])
              for s in W.ORDER))
    check("ask: no relationship in the week has none of its own turns shorter than its filing (every second figure "
          "is a count the device can move)", all(truth[s][1] > 0 for s in W.ORDER))

    # ------------------------------------------------------------------ separation
    nT = B.turns_2026(copy_without(target, ["ssim27", "fleet27", "ssim26"]), reg)
    check("separation: the call reads no 2027 file and no schedule file (deleting them leaves every turn)",
          len(nT) == len(T) and nT["need"].fillna(-1).equals(T["need"].fillna(-1)))
    check("separation: zero 2026 turns carry a DEI 050 flight number (the ask's device lives only in the 2027 file)",
          not set(T["dep_flight"]) & {r[1] for r in refs})

    # ------------------------------------------------------------------ distractors
    tmp = copy_without(target, W.DISTRACTORS)
    fx2 = GO.compute(tmp)
    check("distractors: deleting %s leaves every golden figure" % [W.FILES[k] for k in W.DISTRACTORS],
          GO.figures(fx2) == F and sorted(meta["distractor_files"]) == sorted(W.FILES[k] for k in W.DISTRACTORS))

    # ------------------------------------------------------------------ the golden against figures.json
    check("golden: figures.json is the golden's own figures", GO.figures(fx) == F)
    wb = xlsx_text(os.path.join(golden, W.GOLDEN["workbook"]))
    check("golden: the workbook carries each verdict and the week's counts",
          all(str(fx["ask"][s][0]) in wb for s in W.ORDER) and "Send back" in wb and "Accept" in wb)

    # ------------------------------------------------------------------ input gates
    rows = {}
    for f in files:
        if f.endswith((".csv", ".ssim")):
            rows[f] = sum(1 for _ in open(os.path.join(target, f), encoding="utf-8")) - (1 if f.endswith(".csv") else 0)
    fmts = {f.rsplit(".", 1)[1] for f in files}
    check("gates: %d files, %d formats %s, largest %s rows" % (len(files), len(fmts), sorted(fmts), max(rows.values())),
          len(files) >= 10 and len(fmts) >= 3 and max(rows.values()) >= 25000)

    # ------------------------------------------------------------------ realism sweep
    blobs = {}
    for root in (target, golden):
        for f in os.listdir(root):
            p = os.path.join(root, f)
            data = open(p, "rb").read()
            if f.endswith(".xlsx"):
                with zipfile.ZipFile(p) as z:
                    data = b"".join(z.read(n) for n in z.namelist())
            blobs[f] = data
    tells = {f: [t for t in (b"openpyxl", b"ReportLab", b"reportlab", b"matplotlib", b"pandas") if t in d]
             for f, d in blobs.items()}
    check("realism: containers scrubbed (no writing library named in any file)", not any(tells.values()),
          {f: v for f, v in tells.items() if v})
    texts = dict(docs)
    texts.update({f: open(os.path.join(target, f), encoding="utf-8").read() for f in files if f.endswith((".csv", ".ssim"))})
    texts.update({f: xlsx_text(os.path.join(target, f)) for f in files if f.endswith(".xlsx")})
    texts["golden"] = wb
    for fn in ("prompt.txt", "submission.md", "metadata.json", "DESIGN_NOTE.md"):
        p = os.path.join(TASK, fn)
        if os.path.exists(p):
            texts[fn] = open(p, encoding="utf-8").read()
    check("realism: no em or en dash in any shipped text, golden, prompt, submission, metadata or design note",
          not any(DASH.search(t) for t in texts.values()), [f for f, t in texts.items() if DASH.search(t)])
    mt = {f: dt.datetime.fromtimestamp(os.path.getmtime(os.path.join(target, f))).date().isoformat() for f in files}
    check("realism: every file saved on or before %s" % AS_OF, all(v <= AS_OF for v in mt.values()), mt)
    aodb = pd.read_csv(os.path.join(target, W.FILES["aodb"]))
    check("realism: the stand log runs 1 June to 30 September 2026 and every event id is unique",
          aodb["event_local_time"].min().startswith("2026-06-01") and aodb["event_local_time"].max().startswith(
              "2026-09-30") and aodb["event_id"].is_unique)
    big = {f: n for f, n in rows.items() if n >= 1000}
    check("realism: no large file with a round thousand of rows", all(n % 1000 for n in big.values()), big)
    check("realism: no stand ever holds two aircraft at once", stands_exclusive(aodb))

    print("%s  %d checks failed" % ("PASS" if not bad else "FAIL", len(bad)))
    return not bad


# ---------------------------------------------------------------- variants read from the pack
def turns_unsubcoded(target, reg):
    """Sable Coast's codes read only where they match the grouping file's two digits exactly."""
    import openpyxl  # noqa: F401
    groups = reg["groups"]
    orig = B._codes_group

    def strict(gr, pairs):
        return orig(gr, [(c, m) for c, m in pairs if c and c in gr])
    B._codes_group = strict
    try:
        return B.turns_2026(target, reg)
    finally:
        B._codes_group = orig


def turns_operator_clock(target, reg):
    """The time each turn had on the operators' own block-in times (ACARS), where they differ by a minute from the
    stand log's."""
    T = B.turns_2026(target, reg)
    T = T[T["on_op"].notna()].copy()
    T["B"] = T["std"] - T["on_op"]
    T["need"] = np.where(T["exact"], T["B"] + T["nr"], np.nan)
    return T


def turns_hyphen_kept(target, reg):
    tmp = tempfile.mkdtemp()
    for f in os.listdir(target):
        shutil.copy(os.path.join(target, f), tmp)
    orig = B.stand_visits

    def raw(t):
        ev = {}
        import csv
        for r in csv.DictReader(open(os.path.join(t, W.FILES["aodb"]), encoding="utf-8")):
            tt = dt.datetime.strptime(r["event_local_time"], "%Y-%m-%d %H:%M")
            ev.setdefault(r["registration"], []).append((tt, int(r["event_id"]), r["event_type"], r["stand"],
                                                         r["callsign"]))
        turns = []
        for regn, E in ev.items():
            E.sort()
            cur = None
            for tt, eid, kind, stand, cs in E:
                if kind == "ON-BLOCK":
                    cur = dict(reg=regn.replace("-", ""), on=tt, stand=stand, arr_cs=cs, towed=False)
                elif kind in ("TOW-OFF", "TOW-ON"):
                    if cur is not None:
                        cur["towed"] = True
                elif kind == "OFF-BLOCK":
                    if cur is not None and B.opday(cur["on"]) == B.opday(tt):
                        turns.append(dict(cur, off=tt, dep_cs=cs, dep_stand=stand))
                    cur = None
        return turns
    B.stand_visits = raw
    try:
        return B.turns_2026(tmp, reg)
    finally:
        B.stand_visits = orig


def turns_no_block_dates(target, reg):
    blocks = [(i, lo, hi, op, dt.date(2014, 1, 1), None) for i, lo, hi, op, f, t in reg["blocks"]]
    return B.turns_2026(target, dict(reg, blocks=blocks))


def ssim26_clean(target):
    legs, refs = SS.parse(os.path.join(target, W.FILES["ssim26"]))
    if refs:
        return False
    mv = {}
    for r in pd.read_csv(os.path.join(target, W.FILES["pellmont"]))["flight"]:
        mv[int(r[2:])] = 1
    for r in pd.read_csv(os.path.join(target, W.FILES["grayling"]))["flt"]:
        mv[int(r[2:])] = 1
    from openpyxl import load_workbook
    for row in list(load_workbook(os.path.join(target, W.FILES["sable"]), read_only=True)["KDY movements"]
                    .iter_rows(values_only=True))[1:]:
        mv[int(row[1][2:])] = 1
    return all(lg["flight"] in mv for lg in legs)


def twins_in(legs, refs):
    """Same-schedule pairs on Kelderby's clock in the peak week where neither leg is a partner's duplicate: two
    aircraft."""
    a, b = W.WEEK27
    D = B.dated([lg for lg in legs if (lg["airline"], lg["flight"], lg["leg"]) not in refs], a, b)
    seen, n = set(), 0
    for x in D:
        k = (x["dep"], x["arr"], x["kdep"], x["karr"], x["eqp"])
        n += k in seen
        seen.add(k)
    return n


def stands_exclusive(aodb):
    occ = {}
    a = aodb.sort_values(["event_local_time", "event_id"])
    for t, kind, stand, regn in zip(a["event_local_time"], a["event_type"], a["stand"], a["registration"]):
        r = regn.replace("-", "")
        if kind in ("ON-BLOCK", "TOW-ON"):
            if occ.get(stand) not in (None, r):
                return False
            occ[stand] = r
        else:
            if occ.get(stand) == r:
                occ[stand] = None
    return True


if __name__ == "__main__":
    ok = run(os.path.join(TASK, "target"), os.path.join(TASK, "golden"))
    sys.exit(0 if ok else 1)
