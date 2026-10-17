"""The pack's documents: the authority's schedule facilitation guideline for the E gates, its note on the schedule file
layout, the affiliate register, Sable Coast's movements workbook and the monthly punctuality workbook (a file the
review does not use)."""
import datetime as dt

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import letter                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import inch                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import pack as P      # noqa: E402
import world as W     # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("LiberationSerif", "Regular"), ("LiberationSerif-Bold", "Bold"), ("LiberationSerif-Italic", "Italic"),
              ("LiberationSans", "Regular"), ("LiberationSans-Bold", "Bold"), ("LiberationSans-Italic", "Italic")):
    try:
        pdfmetrics.getFont(nm)
    except KeyError:
        pdfmetrics.registerFont(TTFont(nm, FONT_DIR + nm.split("-")[0] + "-" + f + ".ttf"))
INK = colors.HexColor("#1c1e21")
GREY = colors.HexColor("#5a5f66")
BLUE = colors.HexColor("#1d3f6e")
BODY = ParagraphStyle("body", fontName="LiberationSans", fontSize=9.8, leading=13.4, textColor=INK, spaceAfter=6)
CELL = ParagraphStyle("cell", fontName="LiberationSans", fontSize=8.2, leading=10.2, textColor=INK)
CELLB = ParagraphStyle("cellb", parent=CELL, fontName="LiberationSans-Bold")
H1 = ParagraphStyle("h1", fontName="LiberationSans-Bold", fontSize=14, leading=18, textColor=BLUE, spaceAfter=2)
H1S = ParagraphStyle("h1s", fontName="LiberationSans", fontSize=9, leading=12, textColor=GREY, spaceAfter=10)
H2 = ParagraphStyle("h2", fontName="LiberationSans-Bold", fontSize=10.6, leading=14, textColor=INK, spaceBefore=8,
                    spaceAfter=3)


def _doc(path, title, footer):
    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(GREY)
        canvas.setFont("LiberationSans", 7.5)
        canvas.drawString(0.9 * inch, 0.6 * inch, footer)
        canvas.drawRightString(letter[0] - 0.9 * inch, 0.6 * inch, "Page %d" % doc.page)
        canvas.restoreState()
    d = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.9 * inch, rightMargin=0.9 * inch, topMargin=0.8 * inch,
                          bottomMargin=0.9 * inch, title=title, author=W.AUTHORITY, subject="", creator="",
                          producer="")
    return d, frame


def _table(rows, widths, head="#e3e9f1"):
    data = [[Paragraph(str(v), CELLB if i == 0 else CELL) for v in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa0a6")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(head)),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 4),
                           ("RIGHTPADDING", (0, 0), (-1, -1), 4)]))
    return t


# ---------------------------------------------------------------- the guideline
GUIDELINE = [
    ("1. Scope", [
        "1.1 The E gates are the fourteen ground-level stands of the regional concourse, E1 to E14, where passengers "
        "walk out to the aircraft. Regional flights that cannot be given an E gate stand use the remote stands R1 to "
        "R6, with bussing.",
        "1.2 Before each season, every carrier whose regional flights use the E gates files the minimum ground time it "
        "will schedule at the E gates for each carrier relationship. The authority facilitates the season's schedules at "
        "the E gates on the ground times it accepts. This guideline says how a filing is judged."]),
    ("2. Terms", [
        "2.1 <b>Carrier relationship.</b> One brand selling the flights, the operator flying them and the aircraft type: "
        "for example, Norvale Connect flights flown by Pellmont Aviation's Embraer 175s. A filing is made for each "
        "relationship. A turn belongs to the relationship of its departing flight.",
        "2.2 <b>Turn.</b> An aircraft's visit to an E gate stand between its arrival and its next departure from "
        "Kelderby when both fall on the same operating day, which at Kelderby runs from 03:00 to 02:59 local time. A "
        "turn begins at on-block on the E gate stand and ends at off-block, including any time on a remote stand after "
        "a tow. An aircraft that stays overnight is not turning.",
        "2.3 <b>On-block and off-block</b> are the times the authority's stand log records.",
        "2.4 <b>The time a turn needs</b> runs from on-block until the aircraft is ready to push back: doors closed, "
        "ground equipment clear and the handler's release given. An aircraft that is ready before its scheduled "
        "departure waits for it. Time a ready aircraft is held (by air traffic control, for the stand or the apron, for "
        "the weather or for connecting passengers) is not time the turn needs.",
        "2.5 <b>Delay codes.</b> The operators report each departure's minutes past its scheduled time by IATA delay "
        "code. The authority groups the codes as aircraft not ready or held for another reason, as "
        "kdy_delay_code_groups.csv sets out; an operator's sub-code letter does not change the group."]),
    ("3. The standard", [
        "3.1 The authority accepts a filed minimum ground time for a carrier relationship when no more than 20 per cent "
        "of the relationship's turns would need longer than it.",
        "3.2 A filing for a summer season is judged on the relationship's turns at the E gates in the peak of the "
        "summer before it: turns with on-block from 15 June to 15 September inclusive.",
        "3.3 Ground times are filed in whole multiples of five minutes. When the authority sends a filing back, it gives "
        "the carrier the shortest ground time it would accept, which the carrier may file in its place."]),
    ("4. Records", [
        "4.1 The authority's stand log records each on-block and off-block at every stand and each tow, with the "
        "aircraft registration, the stand and the flight's ATC callsign (the operator's ICAO designator and the flight "
        "number). It is kept in Kelderby local time.",
        "4.2 The operators send their flight logs for every flight to and from Kelderby: scheduled times, block times "
        "and delay codes. The affiliate register lists the brands, the operators flying for them and the flight number "
        "blocks each operator flies; the regional aircraft register lists the aircraft.",
        "4.3 Season schedules and initial submissions are exchanged in the IATA Standard Schedules Information Manual "
        "layout (Chapter 7); the authority's note kdy_schedule_file_layout.pdf gives the positions it reads."]),
]


def write_guideline(path):
    d, frame = _doc(path, "Schedule Facilitation Guideline SFG-3", "Kelderby Airport Authority  |  SFG-3, issue 3")
    story = [Paragraph("Schedule Facilitation Guideline SFG-3", H1),
             Paragraph("Minimum ground times at the E gates. Kelderby Airport Authority, Airside Operations. Issue 3, "
                       "in force from the winter 2025/26 season.", H1S)]
    for head, paras in GUIDELINE:
        story.append(Paragraph(head, H2))
        for p in paras:
            story.append(Paragraph(p, BODY))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Questions on this guideline go to the Airside Performance team at the Kelderby Airport "
                           "Authority.", H1S))
    d.build(story, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------- the schedule file layout note
LAYOUT = [
    ("Record type 1, header", [["1", "Record type (1)"], ["2-35", "Title of contents"], ["195-200", "Record serial number"]]),
    ("Record type 2, carrier", [
        ["1", "Record type (2)"], ["2", "Time mode: L local time, U UTC"], ["3-5", "Airline designator"],
        ["11-13", "Season"], ["15-21", "Period of schedule validity from (DDMMMYY)"],
        ["22-28", "Period of schedule validity to"], ["29-35", "Creation date"], ["36-64", "Title of data"],
        ["195-200", "Record serial number"]]),
    ("Record type 3, flight leg", [
        ["1", "Record type (3)"], ["3-5", "Airline designator"], ["6-9", "Flight number"],
        ["10-11", "Itinerary variation identifier"], ["12-13", "Leg sequence number"], ["14", "Service type"],
        ["15-21", "Period of operation from (DDMMMYY)"], ["22-28", "Period of operation to"],
        ["29-35", "Days of operation (1 Monday to 7 Sunday, blank where the leg does not operate)"],
        ["36", "Frequency rate (blank: every week)"], ["37-39", "Departure station"],
        ["40-43", "Scheduled time of passenger departure (HHMM)"], ["44-47", "Scheduled time of aircraft departure"],
        ["48-52", "UTC/local time variation, departure station"], ["55-57", "Arrival station"],
        ["58-61", "Scheduled time of aircraft arrival"], ["62-65", "Scheduled time of passenger arrival"],
        ["66-70", "UTC/local time variation, arrival station"], ["73-75", "Aircraft type (IATA code)"],
        ["76-95", "Passenger reservations booking designator"], ["138-140", "Onward airline designator"],
        ["141-144", "Onward flight number"], ["145", "Aircraft rotation layover (days)"],
        ["173-192", "Aircraft configuration/version"], ["193", "Date variation, departure"],
        ["194", "Date variation, arrival"], ["195-200", "Record serial number"]]),
    ("Record type 4, segment data", [
        ["1", "Record type (4)"], ["3-14", "As the flight leg record it follows"],
        ["29-30", "Board point and off point indicators"], ["31-33", "Data element identifier (DEI)"],
        ["34-36", "Board point"], ["37-39", "Off point"], ["40-194", "Data"], ["195-200", "Record serial number"]]),
    ("Record type 5, trailer", [["1", "Record type (5)"], ["3-5", "Airline designator"],
                                ["188-193", "Serial number check reference"], ["194", "End code"],
                                ["195-200", "Record serial number"]]),
]
LAYOUT_TEXT = [
    "The authority receives season schedules and initial submissions as fixed-length 200-character records in the "
    "layout of Chapter 7 of the IATA Standard Schedules Information Manual. This note lists the positions the authority "
    "reads. Positions not listed are kept as received.",
    "Files from several carriers are merged in the order they arrive: each carrier's carrier record (type 2) is followed "
    "by its flight leg records (type 3), each with any segment data records (type 4) the carrier sent for it, and its "
    "trailer (type 5). Each carrier's records keep the time mode it filed in.",
    "A flight leg record covers the dates from its period of operation that fall on its days of operation. The onward "
    "airline, flight number and rotation layover give the next flight the aircraft operates. A flight that stops at "
    "Kelderby and continues under the same number has a leg record for each leg, in leg sequence.",
]


def write_layout(path):
    d, frame = _doc(path, "Schedule file layout", "Kelderby Airport Authority  |  Schedule Facilitation")
    story = [Paragraph("Schedule files: the layout the authority reads", H1),
             Paragraph("Kelderby Airport Authority, Schedule Facilitation. Note SF-02, revised March 2025.", H1S)]
    for p in LAYOUT_TEXT:
        story.append(Paragraph(p, BODY))
    for head, rows in LAYOUT:
        story.append(Paragraph(head, H2))
        story.append(_table([["Positions", "Field"]] + rows, [1.0 * inch, 5.6 * inch]))
    d.build(story, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------- workbooks
def _head(ws, row, fill="DCE6F1"):
    from openpyxl.styles import Font, PatternFill
    for c in ws[row]:
        c.font = Font(name="Calibri", bold=True)
        c.fill = PatternFill("solid", fgColor=fill)


AGREEMENTS = [
    ("Norvale Air", "Norvale Connect", "Pellmont Aviation", "Embraer 175", "Norvale Ground Services",
     "capacity purchase agreement", "2016-05-01"),
    ("Norvale Air", "Norvale Connect", "Grayling Air", "ATR 72-600", "Norvale Ground Services",
     "capacity purchase agreement", "2023-04-01"),
    ("Norvale Air", "Norvale Connect", "Sable Coast Air", "Dash 8-400", "Norvale Ground Services",
     "capacity purchase agreement", "2014-10-26"),
    ("Tellura Airways", "Tellura Express", "Pellmont Aviation", "Embraer 175", "Aerocrest Handling",
     "capacity purchase agreement", "2019-03-31"),
    ("Tellura Airways", "Tellura Express", "Sable Coast Air", "Dash 8-400", "Aerocrest Handling",
     "capacity purchase agreement", "2021-03-28"),
    ("Grayling Air", "Grayling Air", "Grayling Air", "ATR 72-600", "Grayling Air (self-handled)", "own flying",
     "2011-06-01"),
]
BLOCK_NOTES = {3450: "transferred to Sable Coast Air flying", 3800: ""}


def write_affiliates(path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Relationships"
    ws.append(["Brand", "Sold as", "Operator", "Aircraft type", "Ground handler at KDY", "Agreement", "Since"])
    _head(ws, 1)
    for r in AGREEMENTS:
        ws.append(list(r))
    for col, w in zip("ABCDEFG", (16, 16, 18, 14, 26, 26, 12)):
        ws.column_dimensions[col].width = w
    ws2 = wb.create_sheet("Flight number blocks")
    ws2.append(["Brand", "IATA", "First number", "Last number", "Operator", "ICAO", "Effective from", "Effective to"])
    _head(ws2, 1)
    for b, op, lo, hi, f, t in W.BLOCKS:
        ws2.append([W.BRANDS[b], b, lo, hi, W.OPERATORS[op][0], op, (f or dt.date(2014, 10, 26)).isoformat(),
                    t.isoformat() if t else None])
    for col, w in zip("ABCDEFGH", (16, 6, 13, 13, 18, 6, 14, 14)):
        ws2.column_dimensions[col].width = w
    ws3 = wb.create_sheet("Operators")
    ws3.append(["ICAO", "Operator", "Telephony", "Air operator certificate"])
    _head(ws3, 1)
    for k, (name, tel) in W.OPERATORS.items():
        ws3.append([k, name, tel, {"PMV": "AOC 7710", "SBC": "AOC 6248", "GYA": "AOC 8155"}[k]])
    for col, w in zip("ABCD", (8, 20, 12, 22)):
        ws3.column_dimensions[col].width = w
    wb.save(path)


def write_sable(path, rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "KDY movements"
    ws.append(P.SABLE_COLS)
    _head(ws, 1, fill="FCE4D6")
    for r in rows:
        ws.append([None if v == "" else v for v in r])
    for row in ws.iter_rows(min_row=2, max_col=1):
        for c in row:
            c.number_format = "dd/mm/yyyy"
    for col, w in zip("ABCDEFGHIJKLM", (11, 9, 9, 6, 6, 7, 7, 10, 10, 8, 6, 8, 6)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(path)


def write_punctuality(path, x):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly punctuality"
    ws.append(["Kelderby International Airport: regional punctuality by brand, summer 2026"])
    ws.append(["On time: within 15 minutes of the scheduled time. Source: operators' monthly returns."])
    ws.append([])
    ws.append(["Month", "Brand", "Departures", "On-time departures", "On-time departure rate", "Arrivals",
               "On-time arrivals", "On-time arrival rate"])
    _head(ws, 4, fill="E2EFDA")
    p = P.punctuality(x)
    for m in (6, 7, 8, 9):
        for b in ("NV", "TQ", "GY"):
            nd, od = p[(m, b, "DEP")]
            na, oa = p[(m, b, "ARR")]
            ws.append([dt.date(2026, m, 1).strftime("%B %Y"), W.REGIONAL[b], nd, od, round(od / nd, 3), na, oa,
                       round(oa / na, 3)])
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
        row[4].number_format = "0.0%"
        row[7].number_format = "0.0%"
    for col, w in zip("ABCDEFGH", (14, 16, 11, 18, 21, 10, 15, 18)):
        ws.column_dimensions[col].width = w
    wb.save(path)
