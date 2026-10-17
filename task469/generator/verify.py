"""An independent recomputation of the golden from the shipped bundle alone (its own readers, its own turn builder, its
own product-limit estimator and SSIM parser; nothing imported from the generator but the file names), compared with
the golden workbook.

    python3 task469/generator/verify.py
"""
import csv
import datetime as dt
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import numpy as np                 # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import world as W                  # noqa: E402  (file names only)

TASK = os.path.dirname(HERE)
MON = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]


def rd(target, name):
    return list(csv.DictReader(open(os.path.join(target, name), encoding="utf-8")))


def opday(stamp):
    """Kelderby's operating day, 03:00 to 02:59, of a 'YYYY-MM-DD HH:MM' local time."""
    return (dt.datetime.strptime(stamp, "%Y-%m-%d %H:%M") - dt.timedelta(hours=3)).date()


def hm(s):
    return int(s[:2]) * 60 + int(s[-2:])


def recompute(target):
    F = W.FILES
    groups = {r["code"]: r["kda_group"].startswith("aircraft not ready") for r in rd(target, F["delay_codes"])}
    wb = load_workbook(os.path.join(target, F["affiliates"]), read_only=True)
    blocks = [r for r in list(wb["Flight number blocks"].iter_rows(values_only=True))[1:]]
    types = {r["registration"].replace("-", ""): r["aircraft_type"] for r in rd(target, F["aircraft"])}
    filings = {(r["brand"], r["operating_carrier"], r["aircraft_type"]): int(r["minimum_ground_time_min"])
               for r in rd(target, F["filings"])}
    regional = {"Norvale Air": "Norvale Connect", "Tellura Airways": "Tellura Express", "Grayling Air": "Grayling Air"}

    def rel(num, day):
        for brand, iata, lo, hi, op, icao, f, t in blocks:
            if lo <= num <= hi and dt.date.fromisoformat(f) <= day and (not t or day <= dt.date.fromisoformat(t)):
                return regional[brand], op
        raise KeyError(num)

    def split(pairs):
        nr = held = 0
        for c, m in pairs:
            if c:
                if groups[c[:2]]:
                    nr += int(m)
                else:
                    held += int(m)
        return nr, held

    dep = {}
    for r in rd(target, F["pellmont"]):
        if r["dep"] == "KDY":
            t = dt.datetime.strptime(r["std_utc"], "%Y-%m-%d %H:%M") - dt.timedelta(hours=4)
            dep[(int(r["flight"][2:]), t.date())] = (t.hour * 60 + t.minute,) + split(
                [(r["delay_code_1"], r["delay_min_1"]), (r["delay_code_2"], r["delay_min_2"])])
    sw = load_workbook(os.path.join(target, F["sable"]), read_only=True)["KDY movements"]
    for row in list(sw.iter_rows(values_only=True))[1:]:
        if row[3] == "KDY":
            d = row[0].date() if isinstance(row[0], dt.datetime) else row[0]
            dep[(int(row[1][2:]), d)] = (hm(row[5]),) + split([(row[9], row[10]), (row[11], row[12])])
    for r in rd(target, F["grayling"]):
        if r["orig"] == "KDY":
            dep[(int(r["flt"][2:]), dt.date.fromisoformat(r["sector_date"]))] = (hm(r["sched_dep"]),) + split(
                [p.split("-") for p in r["delays"].split()])

    ev = {}
    for r in rd(target, F["aodb"]):
        ev.setdefault(r["registration"].replace("-", ""), []).append(
            (r["event_local_time"], int(r["event_id"]), r["event_type"], r["stand"], r["callsign"]))
    pop = {}
    for regn, E in ev.items():
        E.sort()
        on = None
        for t, i, kind, stand, cs in E:
            if kind == "ON-BLOCK":
                on = (t, stand)
            elif kind == "OFF-BLOCK":
                if on and opday(on[0]) == opday(t) and on[1].startswith("E"):
                    day = dt.date.fromisoformat(on[0][:10])
                    if dt.date(2026, 6, 15) <= day <= dt.date(2026, 9, 15):
                        num = int(cs[3:])
                        std, nr, held = dep[(num, day)]
                        brand, op = rel(num, day)
                        key = (brand, op, types[regn])
                        b = std - hm(on[0][11:])
                        pop.setdefault(key, []).append((b, nr))
                on = None
    out = {}
    for key, v in pop.items():
        g0 = filings[key]
        a = np.array(v, dtype=float)
        exact = a[:, 1] > 0
        y = np.where(exact, a[:, 0] + a[:, 1], a[:, 0])

        def share(g):
            F_ = 1.0
            for u in np.unique(y[exact & (y > g)]):
                F_ *= 1 - ((y == u) & exact).sum() / (y <= u).sum()
            return 1 - F_
        shortest = next(g for g in range(20, 75, 5) if share(g) <= 0.20)
        out[key] = dict(turns=len(v), filed=g0, send_back=share(g0) > 0.20, shortest=shortest)
    return out, filings, rel


def read_ssim(path):
    legs, mode = [], "L"
    for raw in open(path, encoding="ascii"):
        s = raw.rstrip("\r\n")
        if s[:1] == "2":
            mode = s[1]
        elif s[:1] == "3":
            legs.append(dict(al=s[2:5].strip(), fl=int(s[5:9]), leg=int(s[11:13]), frm=s[14:21], to=s[21:28],
                             dow=s[28:35], dep=s[36:39], std=s[43:47], arr=s[54:57], sta=s[57:61], eqp=s[72:75],
                             dvar=int(s[47:50]), avar=int(s[65:68]),
                             on=(s[137:140].strip(), int(s[140:144]), int(s[144])) if s[137:140].strip() else None,
                             dv=int(s[193]), mode=mode))
    return legs


def ssim_week(target, rel, filings, types_of):
    """The week's turns from the 2027 submissions, read position by position (aircraft times): each aircraft's chain of
    legs within an operating day, a chain filed twice kept once under the numbers the 2026 season schedule carried."""
    legs = read_ssim(os.path.join(target, W.FILES["ssim27"]))
    op26 = {(x["al"], x["fl"]) for x in read_ssim(os.path.join(target, W.FILES["ssim26"]))}
    lo, hi = dt.date(2027, 7, 10), dt.date(2027, 7, 20)
    D = []
    for x in legs:
        f = dt.date(2000 + int(x["frm"][5:]), MON.index(x["frm"][2:5]) + 1, int(x["frm"][:2]))
        t = dt.date(2000 + int(x["to"][5:]), MON.index(x["to"][2:5]) + 1, int(x["to"][:2]))
        d = max(f, lo)
        while d <= min(t, hi):
            if x["dow"][d.isoweekday() - 1] != " ":
                base = dt.datetime(d.year, d.month, d.day)
                # each end on Kelderby's clock (UTC-4): UTC files directly, local files through the record's own
                # UTC variation
                dvar, avar = (0, 0) if x["mode"] == "U" else (x["dvar"], x["avar"])
                kd = base + dt.timedelta(minutes=hm(x["std"]) - 60 * (dvar + 4))
                ka = base + dt.timedelta(days=x["dv"], minutes=hm(x["sta"]) - 60 * (avar + 4))
                D.append(dict(x, date=d, kd=kd, ka=ka))
            d += dt.timedelta(days=1)
    od = lambda t: (t - dt.timedelta(hours=3)).date()
    idx = {}
    for i, y in enumerate(D):
        idx.setdefault((y["al"], y["fl"], y["date"]), []).append(i)

    def nxt(i):
        y = D[i]
        if y["on"]:
            oa, of, lay = y["on"]
            c = [j for j in idx.get((oa, of, y["date"] + dt.timedelta(days=y["dv"] + lay)), []) if D[j]["dep"] == y["arr"]]
        else:
            c = [j for j in idx.get((y["al"], y["fl"], y["date"]), []) if D[j]["leg"] == y["leg"] + 1]
        return c[0] if c else None
    succ = {}
    for i in range(len(D)):
        j = nxt(i)
        if j is not None and od(D[j]["kd"]) == od(D[i]["kd"]):
            succ[i] = j
    starts = set(range(len(D))) - set(succ.values())
    sigs = {}
    for i in starts:
        chain = [i]
        while chain[-1] in succ:
            chain.append(succ[chain[-1]])
        sig = tuple((D[j]["dep"], D[j]["arr"], D[j]["kd"], D[j]["ka"], D[j]["eqp"]) for j in chain)
        sigs.setdefault(sig, []).append(chain)
    keep = set()
    for chains in sigs.values():
        if len(chains) > 1:
            op = [c for c in chains if all((D[j]["al"], D[j]["fl"]) in op26 for j in c)]
            chains = op if len(op) == 1 else chains
        for c in chains:
            keep.update(c)
    count = {}
    for i in keep:
        y = D[i]
        if y["arr"] != "KDY" or i not in succ or succ[i] not in keep:
            continue
        n = D[succ[i]]
        if not (dt.date(2027, 7, 12) <= y["ka"].date() <= dt.date(2027, 7, 18)):
            continue
        brand, op = rel(n["fl"], y["ka"].date())
        key = (brand, op, types_of[n["eqp"]])
        gt = (n["kd"] - y["ka"]).total_seconds() / 60
        c = count.setdefault(key, [0, 0])
        c[0] += 1
        c[1] += int(gt < filings[key])
    return count


def main(target, golden):
    out, filings, rel = recompute(target)
    names = {"E75": "Embraer 175", "DH4": "Dash 8-400", "AT7": "ATR 72-600"}
    week = ssim_week(target, rel, filings, names)
    wb = load_workbook(os.path.join(golden, W.GOLDEN["workbook"]), read_only=True)
    rows = [r for r in wb["Review"].iter_rows(values_only=True) if r and r[0] in
            ("Norvale Connect", "Tellura Express", "Grayling Air") and isinstance(r[3], int)]
    wrows = [r for r in wb["Summer 2027 peak week"].iter_rows(values_only=True) if r and r[0] in
             ("Norvale Connect", "Tellura Express", "Grayling Air") and isinstance(r[4], int)]
    bad = []
    for r in rows:
        key = (r[0], r[1], r[2])
        v = out[key]
        if (v["filed"], v["turns"], "Send back" if v["send_back"] else "Accept", v["shortest"]) != (r[3], r[4], r[7], r[8]):
            bad.append((key, v, r))
    for r in wrows:
        key = (r[0], r[1], r[2])
        if week.get(key, [0, 0]) != [r[4], r[5]]:
            bad.append((key, week.get(key), r[4:6]))
    ok = not bad and len(rows) == 6 and len(wrows) == 6
    print("%s  verify.py: %d relationships recomputed from the bundle (turns, decision, shortest ground time) and the "
          "2027 week's turns, against the golden workbook%s" % ("PASS" if ok else "FAIL", len(rows),
                                                               "" if ok else ": %s" % bad))
    return ok


if __name__ == "__main__":
    sys.exit(0 if main(os.path.join(TASK, "target"), os.path.join(TASK, "golden")) else 1)
