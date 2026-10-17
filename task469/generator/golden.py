"""The golden deliverable, computed from target/ alone: the airside performance analyst's review workbook for the
Schedule Facilitation Committee (each relationship's verdict and the shortest ground time the authority would accept,
the curves behind them in a chart, and the summer 2027 peak week's turns); and generator/figures.json, every figure the
write-up states."""
import json
import os

import bundle as B
import methods as M
import world as W

LABEL = {"S1": ("Norvale Connect", "Pellmont Aviation", "Embraer 175"),
         "S2": ("Norvale Connect", "Grayling Air", "ATR 72-600"),
         "S3": ("Norvale Connect", "Sable Coast Air", "Dash 8-400"),
         "S4": ("Tellura Express", "Pellmont Aviation", "Embraer 175"),
         "S5": ("Tellura Express", "Sable Coast Air", "Dash 8-400"),
         "S6": ("Grayling Air", "Grayling Air", "ATR 72-600")}
SHORT = {"S1": "Norvale Connect / Pellmont E175", "S2": "Norvale Connect / Grayling ATR 72",
         "S3": "Norvale Connect / Sable Coast Q400", "S4": "Tellura Express / Pellmont E175",
         "S5": "Tellura Express / Sable Coast Q400", "S6": "Grayling Air ATR 72"}
METHOD = M.km


def name(seg):
    b, o, t = LABEL[seg]
    return "%s flown by %s, %s" % (b, o, t) if b != o else "%s, %s" % (b, t)


def compute(target):
    reg = B.registers(target)
    T = B.turns_2026(target, reg)
    pop = B.population(T)
    rv = M.review(pop, reg["filings"], METHOD)
    curves = {seg: [METHOD(pop[pop["seg"] == seg], g) for g in W.GRID] for seg in W.ORDER}
    others = {k: M.review(pop, reg["filings"], f) for k, f in M.ESTIMATORS.items()}
    exact = {seg: int(pop[(pop["seg"] == seg)]["exact"].sum()) for seg in W.ORDER}
    ask = B.ask_a(target, reg)
    return dict(reg=reg, pop=pop, review=rv, curves=curves, others=others, exact=exact, ask=ask,
                rungs=M.rung_verdicts(pop, reg["filings"]), towed=int(pop["towed"].sum()),
                remote_peak=int((T["peak"] & ~T["e_gate"]).sum()))


def pct(v):
    return "%d" % int(v * 100 + 0.5)


def figures(fx):
    rv = fx["review"]
    back = [s for s in W.ORDER if rv[s]["send_back"]]
    keep = [s for s in W.ORDER if not rv[s]["send_back"]]
    rel = {}
    for s in W.ORDER:
        r = rv[s]
        g = r["shortest"]
        rel[s] = {"name": name(s), "filed": str(r["filed"]), "verdict": "send back" if r["send_back"] else "accept",
                  "shortest": str(g), "share_at_filed": pct(r["share"]),
                  "share_at_shortest": pct(fx["curves"][s][W.GRID.index(g)]),
                  "share_5_below_shortest": pct(fx["curves"][s][W.GRID.index(g) - 1]),
                  "turns_2026": "{:,}".format(r["turns"]), "exact_2026": "{:,}".format(fx["exact"][s]),
                  "week_turns": str(fx["ask"][s][0]), "week_short": str(fx["ask"][s][1])}
    r0 = fx["rungs"]["R0 visible overruns"]
    r1 = fx["rungs"]["R1 ground time as need"]
    return {"call": {"send_back": [name(s) for s in back], "accept": [name(s) for s in keep]},
            "relationships": rel,
            "standard_pct": pct(W.STANDARD),
            "rivals": {"visible_overruns_max_pct": pct(max(v["share"] for v in r0.values())),
                       "ground_time_min_pct": pct(min(v["share"] for v in r1.values()))},
            "peak": [W.PEAK_FROM.isoformat(), W.PEAK_TO.isoformat()],
            "week": [W.WEEK27[0].isoformat(), W.WEEK27[1].isoformat()],
            "stated": {"clauses": ["SFG-3 2.1", "SFG-3 2.2", "SFG-3 2.4", "SFG-3 3.1", "SFG-3 3.3"],
                       "operating_day": ["03:00", "02:59"], "season_2026": "2026", "chart_minutes": ["20", "60"]}}


# ---------------------------------------------------------------- the workbook
COLOURS = {"S1": "C0392B", "S4": "E67E22", "S5": "8E44AD", "S2": "2E86C1", "S3": "17A589", "S6": "7F8C8D"}


def write_workbook(path, fx, fig):
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.chart.layout import Layout, ManualLayout
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    thin = Side(style="thin", color="A6A6A6")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = PatternFill("solid", fgColor="DCE6F1")
    back_fill = PatternFill("solid", fgColor="F8CBAD")
    keep_fill = PatternFill("solid", fgColor="E2EFDA")
    bold = Font(name="Calibri", bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    rv = fx["review"]
    back = [s for s in W.ORDER if rv[s]["send_back"]]
    wb = Workbook()

    def table_head(ws, row, height):
        for c in ws[row]:
            if c.value is not None:
                c.font = bold
                c.fill = head
                c.alignment = wrap
                c.border = box
        ws.row_dimensions[row].height = height

    def landscape(ws, area, rows):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_area = area
        ws.print_title_rows = rows

    ws = wb.active
    ws.title = "Review"
    ws["A1"] = "Three of the six summer 2027 ground times filed for the E gates fall short of SFG-3 and go back"
    ws["A1"].font = Font(name="Calibri", bold=True, size=13)
    ws["A2"] = ("Airside Performance for the Schedule Facilitation Committee, 18 November 2026; prepared 29 October "
                "2026. SFG-3 s.3: a filing is accepted when no more than 20% of the relationship's turns at the E gates "
                "from 15 June to 15 September 2026 would need longer than it.")
    ws["A2"].alignment = wrap
    ws.merge_cells("A2:I2")
    ws.row_dimensions[2].height = 30
    ws.append([])
    ws.append(["Brand", "Operator", "Aircraft type", "Filed ground time (min)", "Peak 2026 turns",
               "Turns that held their departure (need seen in full)", "Share needing longer than the filing",
               "Decision", "Shortest ground time accepted (min)"])
    table_head(ws, 4, 45)
    for s in W.ORDER:
        b, o, t = LABEL[s]
        r = rv[s]
        ws.append([b, o, t, r["filed"], r["turns"], fx["exact"][s], r["share"],
                   "Send back" if r["send_back"] else "Accept", r["shortest"]])
        row = ws[ws.max_row]
        for c in row:
            c.border = box
        row[4].number_format = row[5].number_format = "#,##0"
        row[6].number_format = "0%"
        row[7].fill = back_fill if r["send_back"] else keep_fill
        row[7].font = bold
    tot = ws.max_row + 1
    ws.cell(row=tot, column=1, value="All six relationships").font = bold
    ws.cell(row=tot, column=5, value=sum(rv[s]["turns"] for s in W.ORDER)).number_format = "#,##0"
    ws.cell(row=tot, column=6, value=sum(fx["exact"][s] for s in W.ORDER)).number_format = "#,##0"
    for c in ws[tot]:
        c.border = box
    n0 = tot + 2
    notes = [
        "Sent back: %s. Each carrier may refile the shortest ground time in column I (SFG-3 s.3.3); the other three "
        "filings stand." % "; ".join("%s / %s %s at %d minutes" % (LABEL[s][0], LABEL[s][1], LABEL[s][2],
                                                                      rv[s]["filed"]) for s in back),
        "A turn ready before its scheduled departure waited for it, so the stand log shows only that it needed no "
        "longer than the time it had; the shares are product-limit estimates on that basis (see Notes).",
    ]
    for i, t in enumerate(notes):
        cell = ws.cell(row=n0 + i, column=1, value=t)
        cell.alignment = wrap
        ws.merge_cells(start_row=n0 + i, start_column=1, end_row=n0 + i, end_column=9)
        ws.row_dimensions[n0 + i].height = 32
    for col, w in zip("ABCDEFGHI", (17, 18, 14, 12, 11, 17, 14, 11, 15)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    landscape(ws, "A1:I%d" % (n0 + len(notes) - 1), "4:4")

    cs = wb.create_sheet("Need curves")
    cs["A1"] = ("Pellmont's E175s for both brands and Tellura's Q400s stay above 20% at the ground times they filed")
    cs["A1"].font = Font(name="Calibri", bold=True, size=12)
    cs.append(["Ground time (min)"] + [SHORT[s] for s in W.ORDER] + ["SFG-3 standard (20%)"])
    table_head(cs, 2, 45)
    for i, g in enumerate(W.GRID):
        if g > 60:
            break
        cs.append([g] + [fx["curves"][s][i] for s in W.ORDER] + [W.STANDARD])
        for c in cs[cs.max_row][1:]:
            c.number_format = "0%"
    last = cs.max_row
    cs.cell(row=last + 1, column=1, value="Share of each relationship's turns at the E gates, 15 June to 15 September "
                                          "2026, that would need longer than the ground time; product-limit estimates "
                                          "(Notes).").alignment = wrap
    cs.merge_cells(start_row=last + 1, start_column=1, end_row=last + 1, end_column=8)
    cs.row_dimensions[last + 1].height = 30
    ch = LineChart()
    ch.title = "2026 peak turns needing longer than the ground time"
    ch.y_axis.title = "Share of turns"
    ch.x_axis.title = "Ground time (minutes)"
    ch.y_axis.number_format = "0%"
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = 1
    ch.y_axis.majorUnit = 0.2
    ch.legend.position = "r"
    ch.height, ch.width = 11, 24
    ch.add_data(Reference(cs, min_col=2, max_col=2 + len(W.ORDER), min_row=2, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(cs, min_col=1, min_row=3, max_row=last))
    for k, s in enumerate(W.ORDER):
        ln = ch.series[k].graphicalProperties.line
        ln.solidFill = COLOURS[s]
        ln.width = 28000 if s in back else 19000
        ch.series[k].smooth = False
    std = ch.series[-1].graphicalProperties.line
    std.dashStyle = "dash"
    std.solidFill = "404040"
    std.width = 15000
    ch.series[-1].smooth = False
    cs.add_chart(ch, "A%d" % (last + 3))
    for col, w in zip("ABCDEFGH", (11, 15, 15, 15, 15, 15, 15, 13)):
        cs.column_dimensions[col].width = w
    cs.freeze_panes = "B3"
    landscape(cs, "A1:H%d" % (last + 32), "2:2")

    ks = wb.create_sheet("Summer 2027 peak week")
    ks["A1"] = "Every relationship's July 2027 submission already schedules turns under its own filed ground time"
    ks["A1"].font = Font(name="Calibri", bold=True, size=12)
    ks["A2"] = ("Turns the summer 2027 initial submissions schedule at Kelderby, Monday 12 to Sunday 18 July 2027. "
                "Each aircraft once: where a brand files a partner's rotation again under its own numbers (the same "
                "legs, times and type, numbers the 2026 season schedule never carried), the turns are the operating "
                "flights'. A turn is an arrival and the same aircraft's next departure on the same operating day; "
                "through stops included, night stops not.")
    ks["A2"].alignment = wrap
    ks.merge_cells("A2:F2")
    ks.row_dimensions[2].height = 58
    ks.append([])
    ks.append(["Brand", "Operator", "Aircraft type", "Filed ground time (min)", "Turns scheduled",
               "Of which shorter than the filed ground time"])
    table_head(ks, 4, 30)
    for s in W.ORDER:
        b, o, t = LABEL[s]
        ks.append([b, o, t, rv[s]["filed"], fx["ask"][s][0], fx["ask"][s][1]])
        for c in ks[ks.max_row]:
            c.border = box
    tot = ks.max_row + 1
    ks.cell(row=tot, column=1, value="All six relationships").font = bold
    ks.cell(row=tot, column=5, value=sum(fx["ask"][s][0] for s in W.ORDER))
    ks.cell(row=tot, column=6, value=sum(fx["ask"][s][1] for s in W.ORDER))
    for c in ks[tot]:
        c.border = box
    ks.cell(row=tot + 2, column=1, value="The turns under the filing are the through stops, which each carrier "
                                         "schedules shorter than its filed minimum.").alignment = Alignment(
        vertical="top")
    for col, w in zip("ABCDEF", (17, 18, 14, 12, 12, 20)):
        ks.column_dimensions[col].width = w
    ks.freeze_panes = "A5"
    landscape(ks, "A1:F%d" % (tot + 2), "4:4")

    ns = wb.create_sheet("Notes")
    ns["A1"] = "Notes"
    ns["A1"].font = Font(name="Calibri", bold=True, size=12)
    rs = fx["remote_peak"]
    lines = [
        ("Standard", "SFG-3 issue 3, s.3.1 to 3.3: no more than 20% of a relationship's turns would need longer than "
                     "its filing, judged on its turns at the E gates with on-block from 15 June to 15 September 2026; "
                     "ground times in multiples of five minutes."),
        ("Turn", "An aircraft's on-block at an E gate stand and its next off-block on the same operating day, 03:00 "
                 "to 02:59 (s.2.2). %d towed turns are kept to their off-block on the remote stand; %d turns that "
                 "began on a remote stand are outside s.3.2." % (fx["towed"], rs)),
        ("Relationship", "The departing flight's: the callsign's number in the affiliate register's blocks and the "
                         "registration's type in the regional aircraft register (s.2.1)."),
        ("Need", "From on-block to ready for pushback (s.2.4). Where the turn held its departure, the time it had "
                 "(scheduled departure less on-block) plus the minutes coded aircraft not ready; minutes a ready "
                 "aircraft was held (ATC, stand, weather, connections) are left out. Where it was ready in time, the "
                 "need is known only to be no longer than the time it had."),
        ("Estimate", "Under the 2026 schedules most turns had more time than the carriers now file, so their need "
                     "above the new ground times was not recorded. The shares are product-limit estimates for needs "
                     "censored in this way; the share among the turns that had less than the ground time, each of "
                     "which shows whether it needed longer, gives the same decisions and shortest ground times."),
        ("Delay codes", "Grouped as kdy_delay_code_groups.csv sets out (s.2.5); Sable Coast's AHM 731 letters "
                        "dropped. Pellmont's log is in UTC, read at UTC-4."),
        ("Sources", "Stand log export of 2 October 2026; Pellmont, Sable Coast and Grayling flight logs; affiliate "
                    "register (July 2026); regional aircraft register; filings and initial submissions received "
                    "1 to 9 October 2026."),
    ]
    for i, (k, v) in enumerate(lines):
        ns.cell(row=3 + i, column=1, value=k).font = bold
        c = ns.cell(row=3 + i, column=2, value=v)
        c.alignment = wrap
        ns.cell(row=3 + i, column=1).alignment = Alignment(vertical="top")
        ns.row_dimensions[3 + i].height = 15 * (1 + len(v) // 95)
    ns.column_dimensions["A"].width = 14
    ns.column_dimensions["B"].width = 100
    landscape(ns, "A1:B%d" % (2 + len(lines)), "1:1")
    wb.save(path)


def main(target, gdir, fig_path):
    os.makedirs(gdir, exist_ok=True)
    fx = compute(target)
    fig = figures(fx)
    json.dump(fig, open(fig_path, "w", encoding="utf-8"), indent=1)
    write_workbook(os.path.join(gdir, W.GOLDEN["workbook"]), fx, fig)
    return fx, fig
