"""Read the shipped pack, and only the pack, into the tables the review works from: the 2026 turns at Kelderby (the
stand log's visits joined to the operators' logs, the affiliate register and the aircraft register) and the summer
2027 initial submissions' flight legs. Used by the golden, the checks and the verifier alike."""
import csv
import datetime as dt
import os

import numpy as np
import pandas as pd

import ssim as SS
import world as W

F = W.FILES


def _csv(target, key):
    with open(os.path.join(target, F[key]), encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def _min(s):
    h, m = s.split(":")
    return int(h) * 60 + int(m)


def registers(target):
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(target, F["affiliates"]), read_only=True)
    blocks = []
    for r in list(wb["Flight number blocks"].iter_rows(values_only=True))[1:]:
        brand, iata, lo, hi, op, icao, f, t = r
        blocks.append((iata, int(lo), int(hi), icao, dt.date.fromisoformat(f), dt.date.fromisoformat(t) if t else None))
    aircraft = {r["registration"].replace("-", ""): r for r in _csv(target, "aircraft")}
    groups = {r["code"]: r["kda_group"] for r in _csv(target, "delay_codes")}
    filings = {}
    names = {v: k for k, v in W.TYPES.items()}
    ops = {v[0]: k for k, v in W.OPERATORS.items()}
    brands = {v: k for k, v in W.REGIONAL.items()}
    for r in _csv(target, "filings"):
        seg = W.segment_of(brands[r["brand"]], ops[r["operating_carrier"]], names[r["aircraft_type"]])
        filings[seg] = int(r["minimum_ground_time_min"])
    return dict(blocks=blocks, aircraft=aircraft, groups=groups, filings=filings, types=names, ops=ops)


def block_of(blocks, number, day):
    """(brand, operator) of a flight number on a day, from the affiliate register's blocks."""
    for iata, lo, hi, icao, f, t in blocks:
        if lo <= number <= hi and f <= day and (t is None or day <= t):
            return iata, icao
    return None, None


def _since(day, t):
    """Minutes after the given day's midnight."""
    return int((t - dt.datetime(day.year, day.month, day.day)).total_seconds() // 60)


def _codes_group(groups, pairs):
    nr = held = 0
    for code, mins in pairs:
        if not code:
            continue
        g = groups[code[:2]]
        if g == "aircraft not ready":
            nr += int(mins)
        else:
            held += int(mins)
    return nr, held


# ---------------------------------------------------------------- the operators' logs, one row per Kelderby movement
def operator_movements(target, reg):
    """Every movement at Kelderby the three operators logged, on Kelderby's clock: kind, number, date, scheduled time,
    not-ready and held minutes (departures), and the operator's own block time at Kelderby."""
    groups = reg["groups"]
    st = {r["iata"]: int(r["utc_offset_summer"].split(":")[0]) for r in _csv(target, "stations")}
    out = []
    for r in _csv(target, "pellmont"):
        num = int(r["flight"][2:])
        if r["dep"] == W.KDY:
            t = dt.datetime.strptime(r["std_utc"], "%Y-%m-%d %H:%M") + dt.timedelta(hours=W.UTC_OFF)
            a = dt.datetime.strptime(r["out_utc"], "%Y-%m-%d %H:%M") + dt.timedelta(hours=W.UTC_OFF)
            pairs = [(r["delay_code_1"], r["delay_min_1"]), (r["delay_code_2"], r["delay_min_2"])]
            nr, held = _codes_group(groups, pairs)
            out.append(("DEP", num, t.date(), t.hour * 60 + t.minute, nr, held, _since(t.date(), a)))
        else:
            t = dt.datetime.strptime(r["sta_utc"], "%Y-%m-%d %H:%M") + dt.timedelta(hours=W.UTC_OFF)
            a = dt.datetime.strptime(r["in_utc"], "%Y-%m-%d %H:%M") + dt.timedelta(hours=W.UTC_OFF)
            out.append(("ARR", num, t.date(), t.hour * 60 + t.minute, 0, 0, _since(t.date(), a)))
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(target, F["sable"]), read_only=True)
    for row in list(wb["KDY movements"].iter_rows(values_only=True))[1:]:
        d, flight, regn, frm, to, std, sta, off, on, c1, m1, c2, m2 = row
        d = d.date() if isinstance(d, dt.datetime) else d
        num = int(flight[2:])
        if frm == W.KDY:
            nr, held = _codes_group(groups, [(c1, m1), (c2, m2)])
            out.append(("DEP", num, d, _min(std), nr, held, _min(off) + (1440 if _min(off) < _min(std) - 600 else 0)))
        else:
            # scheduled arrival on Kelderby's clock: the leg left on d at the station's time
            dep = _min(std) + 60 * (st[W.KDY] - st[frm])
            arr = _min(sta)
            day = d + dt.timedelta(days=1) if arr < dep - 60 else d
            out.append(("ARR", num, day, arr, 0, 0, _min(on)))
    for r in _csv(target, "grayling"):
        num = int(r["flt"][2:])
        d = dt.date.fromisoformat(r["sector_date"])
        if r["orig"] == W.KDY:
            pairs = [tuple(p.split("-")) for p in r["delays"].split()]
            nr, held = _codes_group(groups, pairs)
            out.append(("DEP", num, d, _min(r["sched_dep"]), nr, held, _min(r["act_dep"])))
        else:
            dep = _min(r["sched_dep"]) + 60 * (st[W.KDY] - st[r["orig"]])
            arr = _min(r["sched_arr"])
            day = d + dt.timedelta(days=1) if arr < dep - 60 else d
            out.append(("ARR", num, day, arr, 0, 0, _min(r["act_arr"])))
    return out


# ---------------------------------------------------------------- the stand log's turns
def stand_visits(target):
    """Each aircraft's visits from the stand log: on-block (with its stand and callsign) to the next off-block, tows
    in between; a visit is a turn when both ends fall on the same operating day."""
    ev = {}
    for r in _csv(target, "aodb"):
        t = dt.datetime.strptime(r["event_local_time"], "%Y-%m-%d %H:%M")
        ev.setdefault(r["registration"].replace("-", ""), []).append(
            (t, int(r["event_id"]), r["event_type"], r["stand"], r["callsign"]))
    turns = []
    for regn, E in ev.items():
        E.sort()
        cur = None
        for t, eid, kind, stand, cs in E:
            if kind == "ON-BLOCK":
                cur = dict(reg=regn, on=t, stand=stand, arr_cs=cs, towed=False)
            elif kind in ("TOW-OFF", "TOW-ON"):
                if cur is not None:
                    cur["towed"] = True
            elif kind == "OFF-BLOCK":
                if cur is not None and opday(cur["on"]) == opday(t):
                    turns.append(dict(cur, off=t, dep_cs=cs, dep_stand=stand))
                cur = None
    return turns


def opday(t):
    """The operating day a local time falls on: 03:00 to 02:59 (SFG-3 2.2)."""
    return (t - dt.timedelta(hours=3)).date()


def turns_2026(target, reg=None):
    reg = reg or registers(target)
    mv = operator_movements(target, reg)
    dep = {(n, d): (s, nr, held, a) for k, n, d, s, nr, held, a in mv if k == "DEP"}
    arr = {(n, d): (s, a) for k, n, d, s, nr, held, a in mv if k == "ARR"}
    rows = []
    for v in stand_visits(target):
        d = v["on"].date()
        dn = int(v["dep_cs"][3:])
        an = int(v["arr_cs"][3:])
        brand, op = block_of(reg["blocks"], dn, d)
        typ = reg["types"][reg["aircraft"][v["reg"]]["aircraft_type"]]
        seg = W.segment_of(brand, op, typ)
        std, nr, held, off_op = dep[(dn, d)]
        on = v["on"].hour * 60 + v["on"].minute
        off = v["off"].hour * 60 + v["off"].minute
        sta, on_op = arr.get((an, d), (None, None))
        B = std - on
        rows.append(dict(seg=seg, day=d, reg=v["reg"], stand=v["stand"], e_gate=v["stand"].startswith("E"),
                         towed=v["towed"], arr_flight=an, dep_flight=dn, through=an == dn, on=on, off=off, std=std,
                         on_op=on_op, off_op=off_op,
                         sta=sta, B=B, nr=nr, held=held, exact=nr > 0, need=B + nr if nr > 0 else np.nan,
                         sched_gt=None if sta is None else std - sta, peak=W.PEAK_FROM <= d <= W.PEAK_TO))
    return pd.DataFrame(rows)


def population(T):
    """The guideline's population: turns at the E gates with on-block in the summer peak."""
    return T[T["e_gate"] & T["peak"]].copy()


# ---------------------------------------------------------------- the 2027 initial submissions
def legs_2027(target):
    legs, refs = SS.parse(os.path.join(target, F["ssim27"]))
    return legs, refs


def dated(legs, lo, hi, utc_as_local=False):
    """Dated legs with each Kelderby end on Kelderby's local clock."""
    out = []
    for lg in legs:
        frm, to = SS.to_date(lg["frm"]), SS.to_date(lg["to"])
        d = max(frm, lo)
        while d <= min(to, hi):
            if lg["dow"][d.isoweekday() - 1] != " ":
                base = dt.datetime(d.year, d.month, d.day)
                dep = base + dt.timedelta(minutes=_min(lg["std"][:2] + ":" + lg["std"][2:]))
                arr = base + dt.timedelta(days=lg["arr_dayvar"], minutes=_min(lg["sta"][:2] + ":" + lg["sta"][2:]))
                if lg["mode"] == "U" and not utc_as_local:
                    dep, arr = dep + dt.timedelta(hours=W.UTC_OFF), arr + dt.timedelta(hours=W.UTC_OFF)
                out.append(dict(lg, date=d, kdep=dep, karr=arr))
            d += dt.timedelta(days=1)
    return out


def turns_2027(legs, refs, reg, dedup="dei", night="clock", through=True, utc_as_local=False):
    """Scheduled turns at Kelderby in the week of 12 July 2027: an arrival and the same aircraft's next departure the
    same local day, followed by the onward link or, at a through stop, the same flight's next leg. The handling
    switches give the schedule ask's other stops."""
    L = list(legs)
    if dedup == "dei":
        L = [x for x in L if (x["airline"], x["flight"], x["leg"]) not in refs]
    a, b = W.WEEK27
    D = dated(L, a - dt.timedelta(days=2), b + dt.timedelta(days=2), utc_as_local=utc_as_local)
    if dedup == "schedule":
        seen, keep = set(), []
        for x in D:
            k = (x["dep"], x["arr"], x["kdep"], x["karr"], x["eqp"])
            if k not in seen:
                seen.add(k)
                keep.append(x)
        D = keep
    deps = {}
    for x in D:
        if x["dep"] == W.KDY:
            deps.setdefault((x["airline"], x["flight"], x["date"]), []).append(x)
    out = []
    for x in D:
        if x["arr"] != W.KDY:
            continue
        nxt = None
        if x["onward"]:
            oa, of, lay = x["onward"]
            if night == "layover" and lay != 0:
                continue
            c = deps.get((oa, of, x["date"] + dt.timedelta(days=x["arr_dayvar"] + lay)), [])
            nxt = c[0] if c else None
        elif through:
            c = [y for y in deps.get((x["airline"], x["flight"], x["date"]), []) if y["leg"] == x["leg"] + 1]
            nxt = c[0] if c else None
        if nxt is None:
            continue
        if night == "clock" and opday(nxt["kdep"]) != opday(x["karr"]):
            continue
        if not (a <= x["karr"].date() <= b):
            continue
        brand, op = block_of(reg["blocks"], nxt["flight"], x["karr"].date())
        seg = W.segment_of(nxt["airline"] if brand is None else brand, op, nxt["eqp"])
        out.append(dict(seg=seg, day=x["karr"].date(), arr=x["karr"], dep=nxt["kdep"],
                        gt=(nxt["kdep"] - x["karr"]).total_seconds() / 60))
    return out


def ask_a(target, reg=None, **handling):
    reg = reg or registers(target)
    legs, refs = legs_2027(target)
    out = {k: [0, 0] for k in W.ORDER}
    for t in turns_2027(legs, refs, reg, **handling):
        if t["seg"] is None:
            continue
        out[t["seg"]][0] += 1
        out[t["seg"]][1] += int(t["gt"] < reg["filings"][t["seg"]])
    return out
