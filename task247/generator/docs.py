"""task247 documents: the watch's methods note, the near-real-time product guide, the reprocessed record's release note
and the 2023 survey protocol (pdf), the head of reef monitoring's request (eml), the 2027 bulletins (xlsx), the data
dictionary and the readme (md)."""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess
import survey

UNI = P.UNIVERSITY
STATION = P.STATION
CENTRE = "Servicio Regional de Analisis Oceanico"
CENTRE_SHORT = "SRAO"


def _addr(who):
    name, user = who
    return "%s <%s@%s>" % (name, user, P.DOMAIN)


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.6, leading=13.4, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=13.5, leading=17, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=14,
                        spaceBefore=8, spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10.5, textColor=colors.HexColor("#444444"))
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=8.4, leading=10.2)
    return body, h1, h2, small, cell


def _grid(rows, widths, cell):
    t = Table([[Paragraph(str(c), cell) for c in r] for r in rows], colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA0AB")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E3E8EF"))]))
    return t


def _footer(text):
    def foot(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(20 * mm, 10 * mm, text)
        canvas.drawRightString(A4[0] - 20 * mm, 10 * mm, "Page %d" % d.page)
        canvas.restoreState()
    return foot


def _doc(path, title, author, subject):
    return SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                             bottomMargin=18 * mm, title=title, author=author, creator=author, subject=subject,
                             invariant=1)


def _sections(s, sec, body, h2):
    for head, paras in sec:
        s.append(Paragraph(head, h2))
        for p in paras:
            s.append(Paragraph(p, body))


# ------------------------------------------------------------------ the request (eml)
def request(path, F):
    text = f"""From: {_addr(P.HEAD)}
To: {_addr(P.COORDINATOR)}
Cc: {_addr(P.ANALYST)}, {_addr(P.DIVE_LEAD)}
Date: Wed, 3 Nov 2027 08:05:12 -0500
Subject: Lista para el muestreo de noviembre
Message-ID: <20271103080512.3317@{P.DOMAIN}>
MIME-Version: 1.0
Content-Type: text/plain; charset="utf-8"
Content-Transfer-Encoding: 8bit

Fabiola,

The Ixchel is booked from 15 November for the post-bleaching survey, and I take the
survey plan to the directorate on the 10th for approval, so I need the site list
before then. As the methods note says, the survey goes to the sites at Alert Level 2
for the season, and the season's levels are set on the reprocessed record, which will
not be out until May.

Everything is in the folder: Rocío's bulletins for the year, the near-real-time
extract to yesterday, the reprocessed record to the end of 2026, the site register
and the climatology. Please give me each site's season maximum as the reprocessed
record will show it, its level and how that compares with the level the bulletins
declared, and a chart I can put in front of the directorate.

For the survey note I also need each site's bleaching prevalence from the 2023
survey, which the report compares against, and the hours each site's logger has read
at or above the site's maximum monthly mean while in the water over the 2024 to 2026
seasons. The loggers in the water now went in in April and come back with the survey.

Alfredo

--
{P.HEAD[0]}
Head of reef monitoring, {STATION}
{UNI}
"""
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(text)


# ------------------------------------------------------------------ the methods note (pdf)
def methods(path, F):
    body, h1, h2, small, cell = _styles()
    doc = _doc(path, "%s: methods note, version 3" % P.WATCH, STATION, "Methods note")
    s = [Paragraph("%s, %s" % (STATION, UNI), small),
         Paragraph("%s: methods note" % P.WATCH, h1),
         Paragraph("Version 3, May 2027. Prepared by the station's reef monitoring group; replaces version 2 (June "
                   "2022).", small), Spacer(1, 4)]
    sec = [
        ("1. Purpose",
         ["The watch follows the heat stress on the station's sixteen reef sites (the site register, %s) through the "
          "heat season, issues a bulletin every week while the season runs, and sets which sites the station's "
          "post-bleaching survey visits in November." % F["sites"]]),
        ("2. The satellite records",
         ["The watch reads two satellite analyses of sea surface temperature (SST) at each site's pixel (the register's "
          "pixel_id), both from the %s (%s) on the same 0.05 degree grid. The near-real-time analysis is issued each "
          "morning for the previous day and is what the weekly bulletins run on (product guide, %s). The reprocessed "
          "analysis is built after each year closes and is released the following May; it is the watch's record of "
          "reference (release note, %s). Both give one value per pixel per day, the day being the UTC day."
          % (CENTRE, CENTRE_SHORT, F["guide"], F["repnote"])]),
        ("3. HotSpots and degree heating weeks",
         ["A site's HotSpot on a day is its SST that day less the site's maximum monthly mean (the register's mmm_c), "
          "when that difference is positive. The maximum monthly mean is the warmest of the twelve monthly means of the "
          "site's daily climatology (%s)." % F["clim"],
          "A site's degree heating weeks (DHW) on a day are the sum of its HotSpots of 1.0 degree Celsius or more over the "
          "84 days ending on that day, the day included, divided by 7, in degree Celsius weeks; smaller HotSpots add "
          "nothing."]),
        ("4. Alert levels and the season",
         ["The watch declares two alert levels: Alert Level 1 when a site's DHW reach 4 degree Celsius weeks and Alert "
          "Level 2 when they reach 8. The heat season runs from 1 June to 31 October. A site's season maximum is the "
          "highest DHW it reaches on a day of the season; the 84 days of a June day's sum reach back into March."]),
        ("5. The weekly bulletins",
         ["Each Tuesday morning the watch issues a bulletin giving, for every site, the SST, HotSpot and DHW of the "
          "latest near-real-time day and the alert level that DHW sets. A site's declared level for the season is the "
          "highest alert level any of the bulletins reporting a day of the season gave it."]),
        ("6. The season's levels and the survey",
         ["A site's level for the season is the alert level its season maximum on the record of reference sets. The "
          "post-bleaching survey goes in November to every site at Alert Level 2 for the season. The record of "
          "reference for a season is released the following May, so the survey list is drawn up in early November "
          "from the station's estimate of each site's season maximum on that record; the survey note gives each "
          "site's estimate, its level and whether the level confirms, raises or lowers the site's declared level."]),
    ]
    _sections(s, sec, body, h2)
    s.append(Paragraph("Document history", h2))
    rows = [["Version", "Date", "Change"],
            ["1", "May 2019", "First issue: monthly heat stress summaries for twelve sites."],
            ["2", "June 2022", "Sixteen sites; weekly bulletins during the season."],
            ["3", "May 2027", "Two alert levels; the survey list and the record of reference (section 6)."]]
    s.append(_grid(rows, [20 * mm, 25 * mm, 125 * mm], cell))
    doc.build(s, onFirstPage=_footer("%s: %s, methods note v3" % (STATION, P.WATCH)),
              onLaterPages=_footer("%s: %s, methods note v3" % (STATION, P.WATCH)))
    postprocess.scrub(path, STATION, "2027-05-18")


# ------------------------------------------------------------------ the near-real-time product guide (pdf)
def guide(path, F):
    body, h1, h2, small, cell = _styles()
    doc = _doc(path, "Near-real-time daily SST analysis, regional 0.05 degree: user guide", CENTRE, "User guide")
    s = [Paragraph("%s (%s)" % (CENTRE, CENTRE_SHORT), small),
         Paragraph("Near-real-time daily sea surface temperature analysis, regional 0.05 degree grid: user guide", h1),
         Paragraph("Product version 3.2. Guide revised July 2027.", small), Spacer(1, 4)]
    sec = [
        ("1. The product",
         ["A daily analysis of sea surface temperature for the western Caribbean on a 0.05 degree grid, blending the "
          "infrared and microwave satellite passes that reach the centre by the processing cut-off. The analysis gives "
          "a value for every ocean pixel on every day. Each day's analysis is issued by about 09:00 UTC on the "
          "following day; the day is the UTC day. Values are in degrees Celsius."]),
        ("2. The observation flag",
         ["Each pixel and day carries an observation flag saying whether a valid satellite observation of the pixel "
          "entered that day's analysis. On a day without one the analysis still provides a value for the pixel, and "
          "the flag marks it as not observed."]),
        ("3. Re-issues",
         ["Observations of a day that reach the centre after its first issue can lead to the day being issued again. "
          "A later issue of a day replaces the earlier one. Each issue carries the issue time and the version of the "
          "processing that produced it."]),
        ("4. Versions",
         ["Version 3.0 ran from the start of the regional product in January 2018. Version 3.1, from the issues of 1 "
          "April 2021, changed the blending weights of the microwave passes; the flag coding did not change. Version "
          "3.2, from the issues of 15 July 2027, moved to the centre's new quality coding, so the flag's values "
          "changed:"]),
    ]
    _sections(s, sec, body, h2)
    rows = [["Flag value", "Versions 3.0 and 3.1", "Version 3.2"],
            ["2", "(not used)", "observed: a valid observation entered the day's analysis"],
            ["1", "observed: a valid observation entered the day's analysis", "not observed"],
            ["0", "not observed", "no analysis (land or ice)"]]
    s.append(_grid(rows, [22 * mm, 70 * mm, 78 * mm], cell))
    sec = [
        ("5. Extracts",
         ["Users who take a site extract receive one row per pixel, day and issue, with the issue time and the version "
          "of each issue. Pixel identifiers follow the centre's regional numbering."]),
        ("6. Accuracy",
         ["The analysis is checked against drifting and moored buoys every month. Differences are smallest where "
          "satellites see the sea every day and larger where cloud keeps them from it; users needing the best estimate "
          "of a past period should use the reprocessed analysis when it is released."]),
    ]
    _sections(s, sec, body, h2)
    doc.build(s, onFirstPage=_footer("%s: near-real-time SST analysis, user guide v3.2" % CENTRE_SHORT),
              onLaterPages=_footer("%s: near-real-time SST analysis, user guide v3.2" % CENTRE_SHORT))
    postprocess.scrub(path, CENTRE, "2027-07-20")


# ------------------------------------------------------------------ the reprocessed record's release note (pdf)
def repnote(path, F):
    body, h1, h2, small, cell = _styles()
    doc = _doc(path, "Reprocessed daily SST analysis, regional 0.05 degree: release note, May 2027", CENTRE,
               "Release note")
    s = [Paragraph("%s (%s)" % (CENTRE, CENTRE_SHORT), small),
         Paragraph("Reprocessed daily sea surface temperature analysis, regional 0.05 degree grid: release note", h1),
         Paragraph("Release of May 2027, covering 1 January 2018 to 31 December 2026.", small), Spacer(1, 4)]
    sec = [
        ("1. The record",
         ["The reprocessed analysis is built after each year closes, on the same grid and pixel numbering as the "
          "near-real-time product. It uses every satellite pass, including the passes that reached the archive too "
          "late for the near-real-time run, and the in situ records from buoys, drifters and floats received after the "
          "fact. It is the centre's record of reference for climate and impact work."]),
        ("2. Releases",
         ["The record is extended once a year: the May release adds the calendar year just ended, and the record for "
          "1 January to 31 December 2027 will be released in May 2028. Values already released are not revised "
          "between releases."]),
        ("3. Format",
         ["One value per pixel per UTC day: analysed_sst, in kelvin, to two decimal places."]),
    ]
    _sections(s, sec, body, h2)
    doc.build(s, onFirstPage=_footer("%s: reprocessed SST analysis, release note May 2027" % CENTRE_SHORT),
              onLaterPages=_footer("%s: reprocessed SST analysis, release note May 2027" % CENTRE_SHORT))
    postprocess.scrub(path, CENTRE, "2027-05-11")


# ------------------------------------------------------------------ the 2023 survey protocol (pdf)
def protocol(path, F):
    body, h1, h2, small, cell = _styles()
    doc = _doc(path, "Post-bleaching survey protocol, 2023", STATION, "Survey protocol")
    s = [Paragraph("%s, %s" % (STATION, UNI), small),
         Paragraph("Post-bleaching survey protocol", h1),
         Paragraph("Issued for the November 2023 survey. Reef monitoring group.", small), Spacer(1, 4)]
    sec = [
        ("1. When and where",
         ["The survey is made in November after a heat season with alert levels. In 2023 it covers all sixteen sites, "
          "after the regional heat of that summer."]),
        ("2. Transects",
         ["At each site four belt transects are laid along the reef, each 20 m long and 1 m wide (0.5 m either side of "
          "the tape), at the site's survey depth as the site register gives it. A transect laid more than 2 m above or "
          "below the survey depth is not part of the site's survey."]),
        ("3. Colonies",
         ["Every stony coral colony of the genera in the annex with a maximum diameter of 10 cm or more whose centre "
          "lies inside the belt is recorded once, on the slate line for the metre of tape its centre lies in (metre 1 "
          "is the tape's first metre), with its genus code, its maximum diameter in centimetres and its condition."]),
        ("4. Condition",
         ["H healthy; P pale; PB partly bleached; BL bleached over the whole colony; D recently dead (bare white "
          "skeleton, no living tissue)."]),
        ("5. Quality assurance",
         ["At each site a second observer repeats transect 1 on a separate sheet marked QA, for the comparison of "
          "observers. QA sheets are not part of the site's figures."]),
        ("6. The site's figures",
         ["A site's bleaching prevalence is the share of the colonies recorded on its transects that are pale, partly "
          "bleached or bleached, in per cent. The report gives each site's prevalence with the number of colonies it "
          "rests on."]),
    ]
    _sections(s, sec, body, h2)
    s.append(Paragraph("Annex. Stony coral genus codes", h2))
    names = {"OFAV": "Orbicella faveolata", "OANN": "Orbicella annularis", "OFRA": "Orbicella franksi",
             "MCAV": "Montastraea cavernosa", "CNAT": "Colpophyllia natans", "DLAB": "Diploria labyrinthiformis",
             "PSTR": "Pseudodiploria strigosa", "SSID": "Siderastrea siderea", "SINT": "Stephanocoenia intersepta",
             "AAGA": "Agaricia agaricites", "PAST": "Porites astreoides", "PPOR": "Porites porites",
             "MMEA": "Meandrina meandrites", "DSTO": "Dichocoenia stokesii", "EFAS": "Eusmilia fastigiata",
             "APAL": "Acropora palmata"}
    rows = [["Code", "Species"]] + [[k, "<i>%s</i>" % names[k]] for k in survey.GENERA]
    s.append(_grid(rows, [25 * mm, 80 * mm], cell))
    doc.build(s, onFirstPage=_footer("%s: post-bleaching survey protocol, 2023" % STATION),
              onLaterPages=_footer("%s: post-bleaching survey protocol, 2023" % STATION))
    postprocess.scrub(path, STATION, "2023-10-24")


# ------------------------------------------------------------------ the bulletins workbook (xlsx)
def bulletins_workbook(path, B):
    wb = Workbook()
    ws = wb.active
    ws.title = "Bulletins"
    ws.cell(row=1, column=1, value="%s, %s: weekly bulletins, 2027" % (STATION, P.WATCH)).font = Font(bold=True)
    ws.cell(row=2, column=1, value="Near-real-time analysis. SST and HotSpot in degrees Celsius, DHW in degree Celsius "
                                   "weeks.")
    head = ["Bulletin", "Issued", "Analysis day", "Site", "Site name", "SST", "HotSpot", "DHW", "Alert level"]
    for j, h in enumerate(head, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE3EA")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    r = 5
    for row in B.itertuples(index=False):
        vals = [row.bulletin, row.issued, row.analysis_day, row.site_code, row.site_name, row.sst_c, row.hotspot_c,
                row.dhw, row.alert_level]
        for j, v in enumerate(vals, 1):
            ws.cell(row=r, column=j, value=v.item() if hasattr(v, "item") else v)
        r += 1
    for k, w in enumerate([9, 12, 13, 7, 22, 8, 9, 7, 14], 1):
        ws.column_dimensions[get_column_letter(k)].width = w
    ws.freeze_panes = "A5"
    wb.save(path)
    postprocess.scrub(path, STATION, "2027-11-02")
    postprocess.finish_office(path, "2027-05-04T13:40:00Z", "2027-11-02T14:05:00Z", UNI)


# ------------------------------------------------------------------ the registers
def dictionary(F):
    return f"""# Data dictionary

## {F["sites"]}
The watch's site register, one row per site: site_code; site_name; region (north, central, south); pixel_id (the
satellite pixel the watch reads for the site); latitude and longitude (decimal degrees, WGS84); survey_depth_m (the
depth of the site's transects and logger, metres); mmm_c (the site's maximum monthly mean, degrees Celsius).

## {F["clim"]}
The daily SST climatology of each pixel: pixel_id; month; day; clim_sst_c (degrees Celsius).

## {F["nrt"]}
The station's extract of the near-real-time analysis, one row per pixel, day and issue: pixel_id; date (the UTC day the
analysis is for, yyyy-mm-dd); sst_c (degrees Celsius); obs_flag (the observation flag, coded as the row's product
version codes it, product guide section 4); product_version; issued_utc (when the issue was made).

## {F["rep"]}
The reprocessed analysis, May 2027 release: pixel_id; time (the UTC day, yyyy-mm-dd); analysed_sst_k (kelvin).

## {F["bulletins"]}
The watch's weekly bulletins for 2027, one row per bulletin and site: the bulletin's number, the date it was issued
(station local date), the analysis day it reports (UTC day), the site, its SST and HotSpot (degrees Celsius), its DHW
(degree Celsius weeks) and the alert level.

## {F["survey"]}
The 2023 post-bleaching survey as entered from the dive slates, one row per slate line: site_code; date; transect (1
to 4); sheet (S survey, QA quality assurance repeat); observer (initials); depth_m (the transect's depth); metre (the
metre of tape the line is written on, 1 to 20); line (the line's order on the sheet); genus (code); max_diameter_cm;
condition (protocol section 4).

## {F["deployments"]}
The dive team's logger log, one row per deployment: site_code; serial; model; depth_m; deployed and retrieved (the date
and time the logger went into and came out of the water, as noted from the boat's GPS on the trip, yyyy-mm-dd hh:mm);
deployed_from and retrieved_from (the boat); notes.

## {F["readings"]}
The station's logger database, every reading downloaded from 1 April 2024: serial; datetime_local (the logger's clock,
station local time, UTC-5); temperature_c (degrees Celsius). A logger records from its launch in the lab to its download.

## {F["kd490"]}
Monthly diffuse attenuation of light at 490 nm from the {CENTRE_SHORT} ocean colour product: pixel_id; month (yyyy-mm);
kd490_per_m (per metre).

## {F["tide"]}
The port tide gauge: datetime_local (station local time); sea_level_m (metres above the gauge's mean sea level).
"""


def readme(F):
    return f"""# {P.WATCH}: November 2027 survey list

Files gathered by the reef monitoring group of the {STATION} ({UNI}) for the November 2027 post-bleaching survey.

- {F["request"]}: the head of reef monitoring's request.
- {F["methods"]}: the watch's methods note (version 3).
- {F["guide"]}: the near-real-time analysis user guide (version 3.2).
- {F["repnote"]}: the reprocessed analysis release note (May 2027).
- {F["sites"]}: the site register.
- {F["clim"]}: the daily climatology.
- {F["nrt"]}: the near-real-time extract, 1 January 2018 to 2 November 2027.
- {F["rep"]}: the reprocessed analysis, 1 January 2018 to 31 December 2026.
- {F["bulletins"]}: the 2027 weekly bulletins.
- {F["protocol"]}: the 2023 post-bleaching survey protocol.
- {F["survey"]}: the 2023 post-bleaching survey.
- {F["deployments"]}: the logger log.
- {F["readings"]}: the logger readings.
- {F["kd490"]}: monthly light attenuation, 2024 to 2027.
- {F["tide"]}: the port tide gauge, 2027.
- DATA_DICTIONARY.md: the fields of each file.
"""


def write_all(target, F, R, B):
    p = lambda k: os.path.join(target, F[k])
    request(p("request"), F)
    methods(p("methods"), F)
    guide(p("guide"), F)
    repnote(p("repnote"), F)
    protocol(p("protocol"), F)
    bulletins_workbook(p("bulletins"), B)
    with open(p("dict"), "w") as fh:
        fh.write(dictionary(F))
    with open(p("readme"), "w") as fh:
        fh.write(readme(F))
