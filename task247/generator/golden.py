"""task247 goldens: the survey workbook (each site's declared level, expected season maximum on the reprocessed record,
level and verdict; the 2023 survey baseline; the loggers' hours at or above the maximum monthly mean) and the note to
the head of reef monitoring with its chart."""
import io
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer

import params as P
import postprocess

XLSX = "reef_heat_stress_2027.xlsx"
PDF = "survey_list_note_2027.pdf"
ORG = "%s, %s" % (P.STATION, P.UNIVERSITY)
AUTHOR = P.COORDINATOR[0]


def _head(ws, row, head, height=46):
    thin = Side(style="thin", color="A6ADB8")
    for j, h in enumerate(head, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(bold=True, size=9.5)
        c.fill = PatternFill("solid", fgColor="DDE4EE")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = Border(bottom=thin)
    ws.row_dimensions[row].height = height


def _workbook(path, A):
    wb = Workbook()
    ws = wb.active
    ws.title = "Survey list"
    ws["A1"] = "2027 heat season: expected levels on the reprocessed record and the November survey list"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Sixteen watch sites. Season maximum degree heating weeks (DHW, degree Celsius weeks), alert levels and "
                "the survey rule as in the watch's methods note v3, sections 3 to 6.")
    ws["A2"].font = Font(italic=True, size=9)
    head = ["Site", "Name", "Region", "Declared level (bulletins)", "Highest DHW in the bulletins",
            "Expected season maximum DHW, reprocessed record", "Level for the season", "Against the declared level",
            "Survey"]
    _head(ws, 4, head, 58)
    for i, s in enumerate(A["sites"], 5):
        vals = [s["code"], s["name"], s["region"], s["declared"], s["declared_dhw"], s["dhw"], s["level"],
                s["verdict"].capitalize(), "Yes" if s["survey"] else "No"]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(size=10, bold=(j == 9 and s["survey"]))
            if j == 5:
                c.number_format = "0.0"
            elif j == 6:
                c.number_format = "0"
            c.alignment = Alignment(horizontal="left" if j in (1, 2, 3, 4, 7, 8) else "right")
            if s["verdict"] == "raises":
                c.fill = PatternFill("solid", fgColor="FBEFD9")
    n = 5 + len(A["sites"])
    notes = [
        "Expected season maximum: each site's near-real-time record (the latest issue of each day) with the clear-sky "
        "warm difference against the reprocessed record (2018 to 2026 observed days) taken out of the observed days and "
        "every other day filled by straight-line interpolation of the anomaly against the climatology between the "
        "observed days either side; HotSpots of 1.0 degree or more above the maximum monthly mean, summed over 84 days "
        "and divided by 7; the highest value on a day from 1 June to 31 October.",
        "Declared level: the highest level the season's bulletins gave the site. Survey: the sites at Alert Level 2 "
        "for the season (section 6).",
        "Sources: SRAO near-real-time extract and reprocessed analysis (May 2027 release), watch site register and "
        "climatology, 2027 weekly bulletins."]
    for k, t in enumerate(notes):
        c = ws.cell(row=n + 1 + k, column=1, value=t)
        c.font = Font(size=8.5, italic=(k == len(notes) - 1))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=n + 1 + k, start_column=1, end_row=n + 1 + k, end_column=len(head))
        ws.row_dimensions[n + 1 + k].height = [58, 18, 18][k]
    for col, w in zip("ABCDEFGHI", [7, 21, 9, 14, 12, 15, 14, 13, 8]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C5"

    sv = wb.create_sheet("2023 survey")
    sv["A1"] = "Bleaching prevalence in the November 2023 post-bleaching survey (baseline for this survey)"
    sv["A1"].font = Font(bold=True, size=11)
    _head(sv, 3, ["Site", "Name", "Bleaching prevalence (%)", "Colonies"], 44)
    for i, s in enumerate(A["sites"], 4):
        for j, v in enumerate([s["code"], s["name"], s["prevalence"], s["colonies"]], 1):
            c = sv.cell(row=i, column=j, value=v)
            c.number_format = "0.0" if j == 3 else ("0" if j == 4 else "General")
            c.alignment = Alignment(horizontal="left" if j <= 2 else "right")
    m = 4 + len(A["sites"])
    c = sv.cell(row=m + 1, column=1, value="Survey protocol 2023: the site's four transects at its survey depth, stony "
                                           "colonies of 10 cm or more, each colony once (a colony of 100 cm or more "
                                           "written on every metre it crosses counted once); prevalence is pale, partly "
                                           "bleached and bleached over all colonies. QA sheets excluded.")
    c.font = Font(size=8.5, italic=True)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    sv.merge_cells(start_row=m + 1, start_column=1, end_row=m + 1, end_column=4)
    sv.row_dimensions[m + 1].height = 44
    for col, w in zip("ABCD", [7, 21, 14, 10]):
        sv.column_dimensions[col].width = w

    bs = wb.create_sheet("Basis")
    bs["A1"] = "Basis of the expected season maxima"
    bs["A1"].font = Font(bold=True, size=11)
    bs["A2"] = ("Near-real-time record to 2 November 2027 (latest issue of each day); reprocessed record, May 2027 release, "
                "1 January 2018 to 31 December 2026. Prepared 8 November 2027.")
    bs["A2"].font = Font(italic=True, size=9)
    bs["A4"] = "Clear-sky difference, near-real-time less reprocessed, observed days 2018 to 2026 (degrees Celsius)"
    bs["B4"] = A["clear_sky"]
    bs["B4"].number_format = "0.00"
    _head(bs, 6, ["Reconstruction of the days the near-real-time analysis did not observe, 2018 to 2026",
                  "Root mean square difference from the reprocessed record (degrees Celsius)",
                  "Mean difference on hot days, anomaly above 0.6 (degrees Celsius)", "Days", "Hot days"], 58)
    names = [("as issued", "As issued (the bulletins' record)"),
             ("flag", "Less the mean differences of observed and filled days"),
             ("regression", "Filled days' anomaly regressed on the reprocessed record's"),
             ("carry", "The last observed anomaly carried"),
             ("nearest", "The nearer observed day's anomaly"),
             ("refill", "Interpolated between the observed days either side (used)")]
    for i, (k, label) in enumerate(names, 7):
        v = A["backtest"][k]
        for j, x in enumerate([label, v["rmse"], v["hot_mean"], v["days"], v["hot_days"]], 1):
            c = bs.cell(row=i, column=j, value=x)
            c.number_format = "0.000" if j in (2, 3) else ("#,##0" if j in (4, 5) else "General")
            c.alignment = Alignment(horizontal="left" if j == 1 else "right")
            c.font = Font(size=10, bold=(k == "refill"))
    for col, w_ in zip("ABCDE", [52, 18, 18, 9, 9]):
        bs.column_dimensions[col].width = w_

    lg = wb.create_sheet("Loggers")
    lg["A1"] = "Hours at or above the maximum monthly mean at the site's logger, June to October 2024 to 2026"
    lg["A1"].font = Font(bold=True, size=11)
    _head(lg, 3, ["Site", "Name", "Hours"], 30)
    for i, s in enumerate(A["sites"], 4):
        for j, v in enumerate([s["code"], s["name"], s["hours"]], 1):
            c = lg.cell(row=i, column=j, value=v)
            c.alignment = Alignment(horizontal="left" if j <= 2 else "right")
    m = 4 + len(A["sites"])
    c = lg.cell(row=m + 1, column=1, value="Hourly readings while the logger was in the water (the log's times, with "
                                           "the Cozumel II's GPS times of August 2025 and April 2026 five hours ahead of "
                                           "local time), station local time.")
    c.font = Font(size=8.5, italic=True)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    lg.merge_cells(start_row=m + 1, start_column=1, end_row=m + 1, end_column=3)
    lg.row_dimensions[m + 1].height = 44
    for col, w in zip("ABC", [7, 21, 10]):
        lg.column_dimensions[col].width = w
    wb.save(path)
    postprocess.scrub(path, AUTHOR, "2027-11-08")
    postprocess.finish_office(path, "2027-11-04T14:20:00Z", "2027-11-08T17:35:00Z", ORG)


def chart_png(A):
    rows = A["sites"]
    fig, ax = plt.subplots(figsize=(6.8, 4.9), dpi=200)
    ax.axvspan(4, 8, color="#FBF1D3", zorder=0, lw=0)
    ax.axvspan(8, 15, color="#F6D9CF", zorder=0, lw=0)
    for x, name in ((4, "Alert Level 1"), (8, "Alert Level 2")):
        ax.axvline(x, color="#8C6D46", lw=0.9, ls="--", zorder=1)
        ax.text(x + 0.1, len(rows) - 0.35, name, fontsize=7.0, color="#6A4F2E", va="center")
    y = list(range(len(rows)))[::-1]
    for yy, r in zip(y, rows):
        col = "#A03A22" if r["verdict"] == "raises" else "#4D6178"
        ax.plot([r["declared_dhw"], r["dhw_exact"]], [yy, yy], color=col, lw=1.4, zorder=2)
        ax.scatter([r["declared_dhw"]], [yy], s=22, facecolor="white", edgecolor=col, lw=1.2, zorder=3)
        ax.scatter([r["dhw_exact"]], [yy], s=26, color=col, zorder=4)
        if r["survey"]:
            ax.text(r["dhw_exact"] + 0.3, yy, "survey", va="center", fontsize=6.6, color=col, fontweight="bold")
    ax.set_yticks(y)
    ax.set_yticklabels(["%s %s" % (r["code"], r["name"]) for r in rows], fontsize=7.0)
    ax.set_xlim(0, 15)
    ax.set_ylim(-0.7, len(rows) - 0.05)
    ax.set_xlabel("Season maximum degree heating weeks, 2027 (degree Celsius weeks)\n(open: declared in the bulletins; "
                  "filled: expected on the reprocessed record; red: level raised)", fontsize=7.0)
    ax.set_title(A["chart_title"], fontsize=9.3, loc="left", fontweight="bold")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.tick_params(axis="x", labelsize=7.5)
    fig.text(0.01, 0.012, "Source: watch bulletins 2027; SRAO near-real-time and reprocessed SST; methods note v3.",
             fontsize=5.8, color="#555555", va="bottom")
    fig.tight_layout(rect=(0, 0.04, 1, 1))
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def _note(path, A):
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=9.2, leading=12.6, spaceAfter=5)
    head = ParagraphStyle("h", parent=body, fontName="Helvetica-Bold", fontSize=12.5, leading=15.5, spaceAfter=3)
    small = ParagraphStyle("s", parent=body, fontSize=7.6, leading=9.8, textColor=colors.HexColor("#444444"))
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=14 * mm,
                            bottomMargin=14 * mm, title=A["title"], author=AUTHOR, creator=ORG,
                            subject="November 2027 post-bleaching survey list", invariant=1)
    s = [Paragraph("%s. To: %s, head of reef monitoring. From: %s, survey coordinator. 8 November 2027."
                   % (ORG, P.HEAD[0], P.COORDINATOR[0]), small),
         Paragraph(A["title"], head)]
    for para in A["paras"]:
        s.append(Paragraph(para, body))
    s += [Image(io.BytesIO(chart_png(A)), width=170 * mm, height=122 * mm), Spacer(1, 2),
          Paragraph(A["footnote"], small)]
    doc.build(s)
    postprocess.scrub(path, AUTHOR, "2027-11-08")


def write(folder, A):
    os.makedirs(folder, exist_ok=True)
    _workbook(os.path.join(folder, XLSX), A)
    _note(os.path.join(folder, PDF), A)
