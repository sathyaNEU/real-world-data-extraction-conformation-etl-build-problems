"""task247 independent verifier: recomputes every graded figure from the shipped bundle alone (no generator module is
imported) and checks figures.json, submission.md and the two goldens against it.

    python3 task247/generator/verify.py [target golden figures.json submission.md]
"""
import csv
import datetime as dt
import json
import os
import sys

import numpy as np
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
F = {"sites": "watch_sites_register.csv", "clim": "srao_sst_daily_climatology_watch_pixels.csv",
     "nrt": "srao_nrt_sst_extract_watch_pixels_2018-01-01_2027-11-02.csv",
     "rep": "srao_reprocessed_sst_watch_pixels_2018_2026.csv", "bulletins": "heat_stress_bulletins_2027.xlsx",
     "survey": "post_bleaching_survey_2023_slates.csv", "deployments": "logger_log_2023-08_2027-04.csv",
     "readings": "logger_readings_from_2024-04-01.csv"}
LEVELS = ["No alert", "Alert Level 1", "Alert Level 2"]
ANNEX = {"OFAV", "OANN", "OFRA", "MCAV", "CNAT", "DLAB", "PSTR", "SSID", "SINT", "AAGA", "PAST", "PPOR", "MMEA", "DSTO",
         "EFAS", "APAL"}


def fail(msg):
    print("VERIFY FAIL:", msg)
    sys.exit(1)


def check(cond, msg):
    if not cond:
        fail(msg)


def rows(target, k):
    with open(os.path.join(target, F[k]), newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def level(x):
    return LEVELS[2] if x >= 8 else (LEVELS[1] if x >= 4 else LEVELS[0])


def series(target):
    sites = rows(target, "sites")
    codes = [s["site_code"] for s in sites]
    pix = {s["pixel_id"]: s["site_code"] for s in sites}
    mmm = {s["site_code"]: float(s["mmm_c"]) for s in sites}
    d0, d1 = dt.date(2018, 1, 1), dt.date(2027, 11, 2)
    days = [d0 + dt.timedelta(i) for i in range((d1 - d0).days + 1)]
    di = {d.isoformat(): i for i, d in enumerate(days)}
    S, T = len(codes), len(days)
    latest = {}
    for r in rows(target, "nrt"):
        key = (r["pixel_id"], r["date"])
        if key not in latest or r["issued_utc"] > latest[key]["issued_utc"]:
            latest[key] = r
    nrt = np.full((S, T), np.nan)
    obs = np.zeros((S, T), bool)
    for (p, d), r in latest.items():
        k, i = codes.index(pix[p]), di[d]
        nrt[k, i] = float(r["sst_c"])
        obs[k, i] = int(r["obs_flag"]) == (2 if r["product_version"] == "3.2" else 1)
    rep = np.full((S, T), np.nan)
    for r in rows(target, "rep"):
        rep[codes.index(pix[r["pixel_id"]]), di[r["time"]]] = float(r["analysed_sst_k"]) - 273.15
    cl = {}
    for r in rows(target, "clim"):
        cl[(r["pixel_id"], int(r["month"]), int(r["day"]))] = float(r["clim_sst_c"])
    inv = {c: p for p, c in pix.items()}
    clim = np.array([[cl[(inv[c], d.month, d.day)] for d in days] for c in codes])
    return codes, sites, mmm, days, di, nrt, obs, rep, clim


def dhw(x, m):
    h = x - m
    h = np.where(h >= 1.0, h, 0.0)
    c = np.concatenate([[0.0], np.cumsum(h)])
    lo = np.maximum(0, np.arange(len(h)) - 83)
    return (c[1:] - c[lo]) / 7.0


def call(target):
    codes, sites, mmm, days, di, nrt, obs, rep, clim = series(target)
    past = np.array([d.year <= 2026 for d in days])
    both = past[None, :] & ~np.isnan(rep) & ~np.isnan(nrt)
    b = float(np.mean((nrt - rep)[both & obs]))
    season = np.array([dt.date(2027, 6, 1) <= d <= dt.date(2027, 10, 31) for d in days])
    idx = np.arange(len(days))
    out = {}
    for k, c in enumerate(codes):
        an = nrt[k] - b - clim[k]
        o = obs[k]
        full = clim[k] + np.where(o, an, np.interp(idx, idx[o], an[o]))
        out[c] = float(dhw(full, mmm[c])[season].max())
    return codes, sites, b, out


def declared(target, codes):
    wb = load_workbook(os.path.join(target, F["bulletins"]), read_only=True, data_only=True)
    ws = wb["Bulletins"]
    best = {c: 0 for c in codes}
    head = None
    for r in ws.iter_rows(values_only=True):
        if r and r[0] == "Bulletin":
            head = list(r)
            continue
        if head is None or r[0] is None:
            continue
        day = str(r[head.index("Analysis day")])[:10]
        if "2027-06-01" <= day <= "2027-10-31":
            c = r[head.index("Site")]
            best[c] = max(best[c], LEVELS.index(r[head.index("Alert level")]))
    return {c: LEVELS[v] for c, v in best.items()}


def survey(target, sites):
    depth = {s["site_code"]: float(s["survey_depth_m"]) for s in sites}
    by = {}
    for r in rows(target, "survey"):
        if r["sheet"] != "S" or r["genus"] not in ANNEX or int(r["max_diameter_cm"]) < 10:
            continue
        if abs(float(r["depth_m"]) - depth[r["site_code"]]) > 2.0:
            continue
        by.setdefault((r["site_code"], r["transect"]), []).append(r)
    n, b = {}, {}
    for (site, tr), lines in by.items():
        lines.sort(key=lambda r: (int(r["metre"]), int(r["line"])))
        prev = set()
        for r in lines:
            key = (r["genus"], r["max_diameter_cm"], r["condition"], int(r["metre"]))
            again = int(r["max_diameter_cm"]) >= 100 and (r["genus"], r["max_diameter_cm"], r["condition"],
                                                          int(r["metre"]) - 1) in prev
            prev.add(key)
            if again:
                continue
            n[site] = n.get(site, 0) + 1
            b[site] = b.get(site, 0) + (r["condition"] in ("P", "PB", "BL"))
    return {c: (100.0 * b[c] / n[c], n[c]) for c in n}


def loggers(target, mmm):
    dep = rows(target, "deployments")
    reads = {}
    for r in rows(target, "readings"):
        reads.setdefault(r["serial"], []).append((dt.datetime.strptime(r["datetime_local"], "%Y-%m-%d %H:%M"),
                                                  float(r["temperature_c"])))
    P = lambda s: dt.datetime.strptime(s, "%Y-%m-%d %H:%M")
    # each trip's clock: the logger's step out of the water against the logged retrieval
    lag = {}
    for r in dep:
        if not r["retrieved"]:
            continue
        z = P(r["retrieved"])
        x = [(t, v) for t, v in reads.get(r["serial"], []) if z - dt.timedelta(hours=9) <= t <= z + dt.timedelta(hours=3)]
        if len(x) < 4:
            continue
        j = max(range(1, len(x)), key=lambda i: x[i][1] - x[i - 1][1])
        lag.setdefault((r["retrieved_from"], r["retrieved"][:7]), []).append((z - x[j][0]).total_seconds() / 3600)
    offset = {k: (5 if np.median(v) > 3 else 0) for k, v in lag.items()}
    hours = {}
    for r in dep:
        if not r["retrieved"]:
            continue
        a = P(r["deployed"]) - dt.timedelta(hours=offset.get((r["deployed_from"], r["deployed"][:7]), 0))
        z = P(r["retrieved"]) - dt.timedelta(hours=offset.get((r["retrieved_from"], r["retrieved"][:7]), 0))
        m = mmm[r["site_code"]]
        hours[r["site_code"]] = hours.get(r["site_code"], 0) + sum(
            1 for t, v in reads.get(r["serial"], []) if a <= t <= z and t.year in (2024, 2025, 2026)
            and 6 <= t.month <= 10 and v >= m)
    return hours, offset


def pdf_text(path):
    from pypdf import PdfReader
    return " ".join(p.extract_text() for p in PdfReader(path).pages).replace("\n", " ")


def main(target, golden, figpath, subpath):
    fig = json.load(open(figpath))
    sub = open(subpath, encoding="utf-8").read()
    codes, sites, b, est = call(target)
    decl = declared(target, codes)
    names = {s["site_code"]: s["site_name"] for s in sites}
    mmm = {s["site_code"]: float(s["mmm_c"]) for s in sites}
    n = 0
    check(abs(round(b, 2) - fig["clear_sky"]) < 1e-9, "the clear-sky difference")
    check("**%.2f degrees Celsius**" % b in sub, "the submission's clear-sky difference")
    on = [c for c in codes if est[c] >= 8]
    check(on == fig["call"] and fig["count"] == len(on), ("the call", on))
    n += 3
    S = {s["code"]: s for s in fig["sites"]}
    for c in codes:
        lv = level(est[c])
        v = "confirms" if lv == decl[c] else ("raises" if LEVELS.index(lv) > LEVELS.index(decl[c]) else "lowers")
        s = S[c]
        check(int(round(est[c])) == s["dhw"] and lv == s["level"] and v == s["verdict"] and decl[c] == s["declared"],
              ("site", c, est[c], lv, v))
        check(0.5 - abs(est[c] - round(est[c])) >= 0.2, ("edge", c, est[c]))
        check("   - %s %s: %d; %s; %s (declared %s)" % (c, names[c], int(round(est[c])), lv, v, decl[c]) in sub,
              ("submission line", c))
        n += 5
    sv = survey(target, sites)
    for c in codes:
        p, k = sv[c]
        check(round(p, 1) == S[c]["prevalence"] and k == S[c]["colonies"], ("ask A", c, p, k))
        check(0.05 - abs(p - round(p, 1)) >= 0.01, ("ask A edge", c, p))
        check("   - %s: %.1f, %d colonies" % (c, round(p, 1), k) in sub, ("ask A line", c))
        n += 3
    hrs, off = loggers(target, mmm)
    check(sorted(k for k, v in off.items() if v) == [("Cozumel II", "2025-08"), ("Cozumel II", "2026-04")], ("offsets", off))
    for c in codes:
        check(hrs[c] == S[c]["hours"], ("ask B", c, hrs[c]))
        check("   - %s: %d" % (c, hrs[c]) in sub, ("ask B line", c))
        n += 2
    # the goldens carry the same figures
    wb = load_workbook(os.path.join(golden, "reef_heat_stress_2027.xlsx"), read_only=True, data_only=True)
    got = {r[0]: r for r in wb["Survey list"].iter_rows(min_row=5, values_only=True) if r[0] in codes}
    check(len(got) == 16, "the golden's 16 rows")
    for c in codes:
        r = got[c]
        check(r[3] == decl[c] and r[5] == int(round(est[c])) and r[6] == level(est[c]) and
              (r[8] == "Yes") == (est[c] >= 8), ("golden row", c))
    a = {r[0]: r for r in wb["2023 survey"].iter_rows(min_row=4, values_only=True) if r[0] in codes}
    h = {r[0]: r for r in wb["Loggers"].iter_rows(min_row=4, values_only=True) if r[0] in codes}
    for c in codes:
        check(abs(a[c][2] - round(sv[c][0], 1)) < 1e-9 and a[c][3] == sv[c][1], ("golden ask A", c))
        check(h[c][2] == hrs[c], ("golden ask B", c))
    n += 48
    # the shipped documents carry the filed sentences the call and the asks lean on
    docs = {"reef_heat_stress_watch_methods_v3.pdf": ["when that difference is positive", "HotSpots of 1.0 degree Celsius or more",
                                                     "84 days ending on that day", "Alert Level 1 when a site's DHW reach 4",
                                                     "runs from 1 June to 31 October", "highest alert level any of the bulletins",
                                                     "every site at Alert Level 2 for the season", "the day being the UTC day"],
            "srao_nrt_sst_user_guide_v3_2.pdf": ["A later issue of a day replaces the earlier one", "from the issues of 15 July 2027"],
            "srao_reprocessed_sst_release_note_2027-05.pdf": ["analysed_sst, in kelvin"],
            "post_bleaching_survey_protocol_2023.pdf": ["is recorded once", "QA sheets are not part of the site's figures",
                                                        "more than 2 m above or below", "10 cm or more",
                                                        "pale, partly bleached or bleached"]}
    for f, phrases in docs.items():
        t = " ".join(pdf_text(os.path.join(target, f)).split())
        for ph in phrases:
            check(ph in t, ("filed sentence", f, ph))
            n += 1
    dct = open(os.path.join(target, "DATA_DICTIONARY.md"), encoding="utf-8").read()
    check("as noted from the boat's GPS" in " ".join(dct.split()) and "station local time, UTC-5" in dct, "the dictionary")
    eml = open(os.path.join(target, "survey_list_request_2027-11-03.eml"), encoding="utf-8").read()
    check("Alert Level 2 for the season" in " ".join(eml.split()), "the request")
    n += 2
    note = pdf_text(os.path.join(golden, "survey_list_note_2027.pdf"))
    for c in on:
        check(c in note, ("the note names", c))
    check("—" not in sub and "–" not in sub, "no dashes")
    n += len(on) + 1
    print("verify: %d checks passed" % n)
    return n


if __name__ == "__main__":
    a = sys.argv[1:]
    if len(a) == 4:
        main(*a)
    else:
        main(os.path.join(TASK, "target"), os.path.join(TASK, "golden"), os.path.join(HERE, "figures.json"),
             os.path.join(TASK, "submission.md"))
