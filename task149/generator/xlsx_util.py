"""Deterministic OOXML writing: openpyxl builds the workbook, then the package is re-zipped with fixed entry times and
the document properties set to the fiction's export date, so two builds give the same bytes and no file carries the
writer's name or the build clock."""
from __future__ import annotations

import datetime as dt
import io
import re
import zipfile

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

STAMP = dt.datetime(2026, 11, 2, 13, 52, 0)
ZIP_TIME = (2026, 11, 2, 13, 52, 0)


class Sheet:
    def __init__(self, name):
        self.name = name
        self.rows = []
        self.widths = []
        self.header_row = 1
        self.number_formats = {}     # column index (1-based) -> format
        self.bold_rows = set()
        self.freeze = None
        self.wrap_cols = set()       # column indices (1-based) whose cells wrap, top-aligned


class Workbook:
    def __init__(self):
        self.sheets = []

    def sheet(self, name):
        s = Sheet(name)
        self.sheets.append(s)
        return s

    def save(self, path, title="", creator="", stamp=STAMP, header=True):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for s in self.sheets:
            ws = wb.create_sheet(s.name)
            for row in s.rows:
                ws.append(row)
            for i, wd in enumerate(s.widths):
                ws.column_dimensions[get_column_letter(i + 1)].width = wd
            if header and s.rows and len(s.rows[0]) > 1 and all(isinstance(x, str) for x in s.rows[0] if x is not None):
                for c in ws[s.header_row]:
                    c.font = Font(bold=True)
            for r_ in s.bold_rows:
                for c in ws[r_]:
                    c.font = Font(bold=True)
            for col, fmt in s.number_formats.items():
                for c in ws[get_column_letter(col)][1:]:
                    c.number_format = fmt
            for col in s.wrap_cols:
                for c in ws[get_column_letter(col)]:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
            if s.freeze:
                ws.freeze_panes = s.freeze
        wb.properties.creator = creator
        wb.properties.lastModifiedBy = creator
        wb.properties.title = title
        wb.properties.created = stamp
        wb.properties.modified = stamp
        buf = io.BytesIO()
        wb.save(buf)
        repack(buf.getvalue(), path, stamp)


def repack(data, path, stamp=STAMP, app="Microsoft Excel"):
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    iso = stamp.strftime("%Y-%m-%dT%H:%M:%SZ")
    for info in zin.infolist():
        body = zin.read(info.filename)
        if info.filename == "docProps/core.xml":
            txt = body.decode("utf-8")
            txt = re.sub(r"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)", r"\g<1>%s\g<2>" % iso, txt)
            txt = re.sub(r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", r"\g<1>%s\g<2>" % iso, txt)
            body = txt.encode("utf-8")
        if info.filename == "docProps/app.xml":
            txt = body.decode("utf-8")
            txt = re.sub(r"<Application>[^<]*</Application>", "<Application>%s</Application>" % app, txt)
            body = txt.encode("utf-8")
        zi = zipfile.ZipInfo(info.filename, date_time=(stamp.year, stamp.month, stamp.day, stamp.hour, stamp.minute, 0))
        zi.compress_type = zipfile.ZIP_DEFLATED
        zi.external_attr = 0o600 << 16
        zout.writestr(zi, body)
    zout.close()
    with open(path, "wb") as f:
        f.write(out.getvalue())
