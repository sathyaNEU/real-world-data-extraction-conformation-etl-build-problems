"""Independent verifier for task149: reads only the shipped pack (target/) and the goldens, never the generator's
state or calc.py, and recomputes on its own code path (csv, json, openpyxl and pyarrow readers; the Poisson fit with a
baseline for each DC by profiling the baselines out and solving the response's score equation with Newton steps; the
day counts walked day by day; the item weights looked up by date) the forecast, the default path, the chart's five
fiscal years and both tables, then checks them against figures.json, the workbook and the slide."""
from __future__ import annotations

import csv
import datetime as dt
import json
import math
import os
import sys
import zoneinfo
from collections import defaultdict

import numpy as np
import openpyxl
import pyarrow.parquet as pq

ZONE_OF_STATE = {"OH": "America/New_York", "IN": "America/Indiana/Indianapolis", "MI": "America/Detroit",
                 "PA": "America/New_York", "KY": "America/New_York", "GA": "America/New_York",
                 "NC": "America/New_York", "IL": "America/Chicago", "TN": "America/Chicago", "TX": "America/Chicago",
                 "WI": "America/Chicago"}
OLD_CODES = {"DRYG": "AMB", "PERI": "CHL", "FROZ": "FRZ", "BULK": "HVB"}
DIRECT = ("RCV", "PUT", "SEL", "LOD")


def last_saturday_of_december(y):
    d = dt.date(y, 12, 31)
    while d.weekday() != 5:
        d -= dt.timedelta(days=1)
    return d


def fiscal_year_start(fy):
    return last_saturday_of_december(fy - 1) + dt.timedelta(days=1)


def fiscal_week(d):
    for fy in range(2021, 2028):
        if fiscal_year_start(fy) <= d < fiscal_year_start(fy + 1):
            return fy, (d - fiscal_year_start(fy)).days // 7 + 1
    return None, None


def fiscal_period(wk):
    ends = [4, 8, 13, 17, 21, 26, 30, 34, 39, 43, 47, 52]
    for i, e in enumerate(ends):
        if wk <= e:
            return i + 1
    return 12


def read_csv(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def sheet_rows(path, name, header_row=1):
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)[name]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[header_row - 1]
    return [dict(zip(hdr, r)) for r in rows[header_row:] if any(v is not None for v in r)]


class Pack:
    def __init__(self, T):
        self.T = T
        reg = read_csv(os.path.join(T, "site_register_2026-11-02.csv"))
        self.canon, self.loc, self.est, self.state, self.prog, self.label = {}, {}, {}, {}, [], {}
        for r in reg:
            nums = [int(x) for x in r["DC_NO"].split(";")]
            for n in nums:
                self.canon[n] = nums[0]
            for l_ in r["WMS_LOC"].split(";"):
                self.loc[l_.strip()] = nums[0]
            self.est[r["EHS_ESTABLISHMENT"]] = nums[0]
            self.state[nums[0]] = r["STATE"]
            self.label[nums[0]] = (str(nums[-1]), r["DC_NAME"])
            if r["ERGO_PROGRAM_FY27"] == "Y":
                self.prog.append(nums[0])
        self.prog.sort()

    def recordables(self):
        """Every recordable case once: (dc, local date of injury, case number, source, osha type, status)."""
        with open(os.path.join(self.T, "ehs_platform_cases_2024-07-01_2026-11-02.json"), encoding="utf-8") as fh:
            plat = json.load(fh)["cases"]
        self.plat = {c["case_id"]: c for c in plat}
        moved = {c["legacy_case_ref"] for c in plat if c["legacy_case_ref"]}
        out = []
        for r in read_csv(os.path.join(self.T, "ehs_legacy_case_export_fy21_2024-06-30.csv")):
            if r["CASE_NO"] in moved or r["CLASSIFICATION"] != "RECORDABLE" or r["CASE_STATUS"] == "NOT WORK RELATED":
                continue
            m, d, y = (int(x) for x in r["INCIDENT_DATE"].split("/"))
            out.append((self.est[r["ESTABLISHMENT"]], dt.date(y, m, d), r["CASE_NO"], "legacy", r["OSHA_CASE_TYPE"],
                        r["CASE_STATUS"]))
        for c in plat:
            if c["classification"]["category"] != "Recordable" or c["status"].startswith("Removed"):
                continue
            dc = self.canon[int(c["site"]["site_code"].replace("DC", ""))]
            t = dt.datetime.fromisoformat(c["occurred_at"].replace("Z", "+00:00"))
            day = t.astimezone(zoneinfo.ZoneInfo(ZONE_OF_STATE[self.state[dc]])).date()
            out.append((dc, day, c["case_id"], "platform", c["classification"]["osha_case_type"], c["status"]))
        return out


def standard_minutes(T):
    path = os.path.join(T, "ie_engineered_labor_standards.xlsx")
    mins = {}
    for r in sheet_rows(path, "Minutes per case", header_row=4):
        if r["Function"] in DIRECT:
            mins[r["Function"]] = {c: float(r[c]) for c in ("AMB", "CHL", "FRZ", "HVB")}
    factors = defaultdict(list)
    for r in sheet_rows(path, "Site factors"):
        a = dt.date.fromisoformat(str(r["EFFECTIVE_FROM"])[:10])
        b = dt.date.fromisoformat(str(r["EFFECTIVE_TO"])[:10]) if r["EFFECTIVE_TO"] else None
        factors[(int(r["DC_NO"]), r["FUNCTION"])].append((a, b, float(r["FACTOR"])))
    return mins, factors


def factor(factors, canon, dc, fn, day):
    for (n, f), spans in factors.items():
        if canon[n] == dc and f == fn:
            for a, b, x in spans:
                if a <= day and (b is None or day <= b):
                    return x
    raise KeyError((dc, fn, day))


def std_hours(cases, mins, fp, fs):
    """Standard hours of a bundle of cases by category at the PUT and SEL site factors."""
    tot = 0.0
    for c, n in cases.items():
        tot += n * (mins["RCV"][c] + mins["LOD"][c] + mins["PUT"][c] * fp + mins["SEL"][c] * fs)
    return tot / 60.0


def poisson_fe(groups, y, x, expo, iters=60):
    """y ~ Poisson(expo * exp(a_group + g * x)): the baselines profiled out, the score for g solved by Newton."""
    groups, y, x, expo = np.asarray(groups), np.asarray(y, float), np.asarray(x, float), np.asarray(expo, float)
    keys, inv = np.unique(groups, return_inverse=True)
    ysum = np.bincount(inv, weights=y)
    g = 0.0
    for _ in range(iters):
        w = expo * np.exp(g * x)
        a = ysum / np.bincount(inv, weights=w)
        mu = w * a[inv]
        s0 = np.bincount(inv, weights=mu)
        s1 = np.bincount(inv, weights=mu * x)
        s2 = np.bincount(inv, weights=mu * x * x)
        score = float(np.sum(x * (y - mu)))
        info = float(np.sum(s2 - s1 * s1 / np.where(s0 > 0, s0, 1)))
        step = score / info
        g += step
        if abs(step) < 1e-12:
            break
    w = expo * np.exp(g * x)
    a = ysum / np.bincount(inv, weights=w)
    return g, dict(zip(keys.tolist(), np.log(a).tolist()))


def main_call(P):
    T = P.T
    rec = P.recordables()
    inj = defaultdict(int)
    for dc, day, *_ in rec:
        fy, wk = fiscal_week(day)
        if fy and (fy, wk) <= (2026, 44):
            inj[(dc, fy, fiscal_period(wk))] += 1
    hours, direct = defaultdict(float), defaultdict(float)
    week_dc = set()
    for r in read_csv(os.path.join(T, "wfm_hours_by_dept_week_fy21_fy26.csv")):
        dc = P.canon[int(r["DC_NO"])]
        fy, wk = int(r["FISCAL_YEAR"]), int(r["FISCAL_WEEK"])
        h = float(r["REG_HRS"]) + float(r["OT_HRS"]) + float(r["AGY_HRS"])
        k = (dc, fy, fiscal_period(wk))
        hours[k] += h
        if r["DEPT"] in DIRECT:
            direct[k] += h
        week_dc.add((dc, fy, wk))
    mins, factors = standard_minutes(T)
    tb = pq.read_table(os.path.join(T, "wms_cases_shipped_by_category_week.parquet")).to_pydict()
    wk_cases = defaultdict(lambda: defaultdict(float))
    for loc, fy, wk, cat, n in zip(tb["loc_code"], tb["fiscal_year"], tb["fiscal_week"], tb["category_code"],
                                   tb["cases_shipped"]):
        wk_cases[(P.loc[loc], int(fy), int(wk))][OLD_CODES.get(cat, cat)] += float(n)
    earned = defaultdict(float)
    fcache = {}
    for (dc, fy, wk), cs in wk_cases.items():
        day = fiscal_year_start(fy) + dt.timedelta(days=7 * (wk - 1))
        key = (dc, day)
        if key not in fcache:
            fcache[key] = (factor(factors, P.canon, dc, "PUT", day), factor(factors, P.canon, dc, "SEL", day))
        fp, fs = fcache[key]
        earned[(dc, fy, fiscal_period(wk))] += std_hours(cs, mins, fp, fs)
    keys = sorted(hours)
    y = [inj.get(k, 0) for k in keys]
    x = [math.log(earned[k] / direct[k]) for k in keys]
    g, base = poisson_fe([k[0] for k in keys], y, x, [hours[k] for k in keys])
    # the plan
    plan = os.path.join(T, "fy2027_operating_plan_v3.xlsx")
    vol = defaultdict(dict)
    for r in sheet_rows(plan, "Volume"):
        vol[(P.canon[int(r["DC_NO"])], int(r["PERIOD"]))][r["CATEGORY"]] = float(r["CASES"])
    lab = {}
    for r in sheet_rows(plan, "Labor hours"):
        lab[(P.canon[int(r["DC_NO"])], int(r["PERIOD"]))] = (float(r["DIRECT_HOURS"]), float(r["INDIRECT_HOURS"]))
    d0 = fiscal_year_start(2027)
    fc, by_dc = 0.0, defaultdict(float)
    l12 = {}
    for dc in P.prog:
        fp, fs = factor(factors, P.canon, dc, "PUT", d0), factor(factors, P.canon, dc, "SEL", d0)
        for per in range(1, 13):
            dh, ih = lab[(dc, per)]
            e = std_hours(vol[(dc, per)], mins, fp, fs)
            mu = (dh + ih) * math.exp(base[dc] + g * math.log(e / dh))
            fc += mu
            by_dc[dc] += mu
    # the default path: the latest twelve months' rate per hour worked times the planned hours
    rec12 = defaultdict(int)
    for dc, day, *_ in rec:
        fy, wk = fiscal_week(day)
        if fy and ((fy == 2025 and wk >= 45) or (fy == 2026 and wk <= 44)):
            rec12[dc] += 1
    h12 = defaultdict(float)
    for r in read_csv(os.path.join(T, "wfm_hours_by_dept_week_fy21_fy26.csv")):
        fy, wk = int(r["FISCAL_YEAR"]), int(r["FISCAL_WEEK"])
        if (fy == 2025 and wk >= 45) or fy == 2026:
            h12[P.canon[int(r["DC_NO"])]] += float(r["REG_HRS"]) + float(r["OT_HRS"]) + float(r["AGY_HRS"])
    default = sum(sum(sum(lab[(dc, p)]) for p in range(1, 13)) * rec12[dc] / h12[dc] for dc in P.prog)
    chart = {fy: sum(1 for dc, day, *_ in rec if dc in P.prog and fiscal_week(day)[0] == fy) for fy in range(2021, 2026)}
    return dict(forecast=fc, gamma=g, default=default, chart=chart, by_dc=dict(by_dc))


def dart_days(P):
    """Each programme DC's fiscal 2026 recordable cases to 31 October 2026: days away and restricted or transferred,
    walked day by day on the work-status log, a transferred associate's continuation file followed."""
    cut = dt.date(2026, 10, 31)
    rows = read_csv(os.path.join(P.T, "rtw_work_status_log_fy2026.csv"))
    files = defaultdict(list)
    for r in rows:
        files[r["RTW_CASE_ID"]].append(r)
    by_case = defaultdict(list)
    cont_start = {}
    for rid, rs in files.items():
        if rs[0]["EHS_CASE_NO"]:
            by_case[rs[0]["EHS_CASE_NO"]].append(rid)
        else:
            first = min(rs, key=lambda z: z["EFFECTIVE_DATE"])
            cont_start[(first["ASSOCIATE_NO"], first["EFFECTIVE_DATE"])] = rid
    out = {dc: 0 for dc in P.prog}
    fy26 = fiscal_year_start(2026)
    for cid, c in P.plat.items():
        if c["classification"]["category"] != "Recordable" or c["status"].startswith("Removed"):
            continue
        dc = P.canon[int(c["site"]["site_code"].replace("DC", ""))]
        if dc not in out:
            continue
        t = dt.datetime.fromisoformat(c["occurred_at"].replace("Z", "+00:00"))
        day0 = t.astimezone(zoneinfo.ZoneInfo(ZONE_OF_STATE[P.state[dc]])).date()
        if not (fy26 <= day0 <= cut) or cid not in by_case:
            continue
        events = []
        for rid in by_case[cid]:
            for r in files[rid]:
                events.append((r["EFFECTIVE_DATE"], r["STATUS"], r["ASSOCIATE_NO"]))
                if r["STATUS"] == "TRANSFER OUT":
                    nxt = cont_start.get((r["ASSOCIATE_NO"], r["EFFECTIVE_DATE"]))
                    if nxt:
                        for z in files[nxt]:
                            events.append((z["EFFECTIVE_DATE"], z["STATUS"], z["ASSOCIATE_NO"]))
        status_on = {}
        rank = {"TRANSFER OUT": 0, "AWAY": 1, "RESTRICTED": 2, "TRANSFER": 2, "FULL DUTY": 3, "SEPARATED": 3}
        for d_, st, _ in sorted(events, key=lambda e: (e[0], rank[e[1]])):
            status_on[dt.date.fromisoformat(d_)] = st
        state, n = None, 0
        day = day0 + dt.timedelta(days=1)
        while day <= cut and n < 180:
            if day in status_on and status_on[day] != "TRANSFER OUT":
                state = status_on[day]
            if state in ("FULL DUTY", "SEPARATED"):
                break
            if state in ("AWAY", "RESTRICTED", "TRANSFER"):
                n += 1
            day += dt.timedelta(days=1)
        out[dc] += n
    return out


def heavy_cases(P):
    """Each programme DC's cases picked over 50 lb, fiscal 2026 period 1 to week 44, each SKU weighed at the version in
    effect during the period it was picked, kilogram weights (vendor portal) in pounds."""
    hist = defaultdict(list)
    for r in read_csv(os.path.join(P.T, "wms_item_master_history.csv")):
        a = dt.date.fromisoformat(r["EFFECTIVE_FROM"])
        b = dt.date.fromisoformat(r["EFFECTIVE_TO"]) if r["EFFECTIVE_TO"] else dt.date(2099, 1, 1)
        w = float(r["CASE_WEIGHT"]) * (2.20462 if r["WEIGHT_SOURCE"] == "PORTAL" else 1.0)
        hist[int(r["SKU"])].append((a, b, w))
    start = fiscal_year_start(2026)
    ends = [4, 8, 13, 17, 21, 26, 30, 34, 39, 43, 44]
    span = {}
    for p in range(1, 12):
        w0 = 1 if p == 1 else ends[p - 2] + 1
        span[p] = (start + dt.timedelta(days=7 * (w0 - 1)), start + dt.timedelta(days=7 * ends[p - 1] - 1))
    tb = pq.read_table(os.path.join(P.T, "wms_sku_picks_fy2026_p01_p11.parquet")).to_pydict()
    out = {dc: 0 for dc in P.prog}
    for loc, per, sku, n in zip(tb["loc_code"], tb["fiscal_period"], tb["sku"], tb["cases_picked"]):
        a, b = span[int(per)]
        ws = {round(w, 6) for f, t, w in hist[int(sku)] if f <= b and t >= a}
        assert len(ws) == 1, (sku, per, ws)
        if ws.pop() > 50.0:
            out[P.loc[loc]] += int(n)
    return out


def num(s):
    return float(str(s).replace(",", ""))


def main(target, golden, figures_path):
    P = Pack(target)
    F = json.load(open(figures_path, encoding="utf-8"))
    bad = []
    M = main_call(P)
    fc = int(round(M["forecast"] / 100.0)) * 100
    if fc != num(F["call"]["forecast_recordable_cases"]):
        bad.append(("forecast", fc, F["call"]["forecast_recordable_cases"]))
    if abs(M["forecast"] - num(F["call"]["unrounded"])) > 0.5:
        bad.append(("forecast unrounded", M["forecast"], F["call"]["unrounded"]))
    if abs(M["gamma"] - num(F["critical_components"]["response_to_log_pace"])) > 0.05:
        bad.append(("response", M["gamma"], F["critical_components"]["response_to_log_pace"]))
    if abs(M["default"] - num(F["default_path"]["unrounded"])) > 0.5:
        bad.append(("default", M["default"], F["default_path"]["unrounded"]))
    for fy, v in M["chart"].items():
        if v != num(F["chart_fiscal_years"][str(fy)]):
            bad.append(("chart", fy, v, F["chart_fiscal_years"][str(fy)]))
    A = dart_days(P)
    for row in F["ask_a_dart_days"]:
        dc = P.canon[int(row["dc"])]
        if A[dc] != num(row["days"]):
            bad.append(("ask A", row["dc"], A[dc], row["days"]))
    Bh = heavy_cases(P)
    for row in F["ask_b_cases_over_50_lb"]:
        dc = P.canon[int(row["dc"])]
        if Bh[dc] != num(row["cases"]):
            bad.append(("ask B", row["dc"], Bh[dc], row["cases"]))
    # the goldens carry the same figures
    wb = openpyxl.load_workbook(os.path.join(golden, "ergonomics_forecast_fy2027.xlsx"), data_only=True)
    if wb["Forecast"]["B3"].value != fc:
        bad.append(("workbook forecast", wb["Forecast"]["B3"].value))
    for name, ref in (("DART days FY2026", A), ("Cases over 50 lb FY2026", Bh)):
        ws = wb[name]
        got = {str(r[0]): r[2] for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0] and r[0] != "Total"
               and r[1]}
        for dc in P.prog:
            if got.get(P.label[dc][0]) != ref[dc]:
                bad.append(("workbook", name, dc, got.get(P.label[dc][0]), ref[dc]))
    from pptx import Presentation
    prs = Presentation(os.path.join(golden, "ergonomics_forecast_fy2027.pptx"))
    slide = prs.slides[0]
    texts = " ".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
    if "{:,}".format(fc) not in texts:
        bad.append(("slide forecast", texts[:120]))
    charts = [sh.chart for sh in slide.shapes if sh.has_chart]
    vals = list(charts[0].plots[0].series[0].values)
    if vals != [float(M["chart"][y]) for y in range(2021, 2026)] + [float(fc)]:
        bad.append(("slide chart", vals))
    if len(prs.slides) != 1:
        bad.append(("slides", len(prs.slides)))
    print("verify.py: forecast %.1f (%d), response %.3f, default %.1f, chart %s; Ask A %d and Ask B %d figures" % (
        M["forecast"], fc, M["gamma"], M["default"], M["chart"], len(A), len(Bh)))
    if bad:
        for b in bad:
            print("VERIFY FAIL", b)
        raise SystemExit(1)
    print("verify.py: every figure agrees with figures.json, the workbook and the slide")


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    task = os.path.dirname(here)
    main(os.path.join(task, "target"), os.path.join(task, "golden"), os.path.join(here, "figures.json"))
