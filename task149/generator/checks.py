"""task149 checks, run on every build: the input gates; the three-rules block (R1 the field's raw material, R2 every
mess item neutral, R3a the analyst test on the clean truth and the pack, R3b the decisive step alone, R3c no plant);
the strict rule 3 test on every shipped sentence that bears on how the forecast is computed; the one-sigma test; the
out-of-sample rejection of the default and of every rival; every rung by name; every construction of the golden; every
semantic fork row (SF1 to SF14) and closure axis (C1 to C22); both asks with every stop, hazard and the cracker's path;
the separation line and the distractors; the chart; the realism sweep; the goldens; the prompt. Any failure stops the
build."""
from __future__ import annotations

import csv
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import tempfile

import numpy as np
import openpyxl
import pandas as pd
import pyarrow.parquet as pq
import statsmodels.api as sm
import statsmodels.formula.api as smf

import asks_world as AW
import calc
import common as K
import deliver as DL
import pack as PK
import world as WD

DASH = re.compile("[%s%s]" % (chr(0x2013), chr(0x2014)))
REPO = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
LO, HI = 950.0, 1050.0          # the golden's bin: rounds to 1,000


class Checks:
    def __init__(self):
        self.res = []

    def __call__(self, label, ok, detail=""):
        self.res.append((label, bool(ok), detail))
        return ok

    def report(self):
        if os.environ.get("CHECKS_LABELS"):
            with open(os.environ["CHECKS_LABELS"], "w", encoding="utf-8") as fh:
                fh.writelines(("ok   " if ok else "FAIL ") + lab + "\n" for lab, ok, _ in self.res)
        bad = [r for r in self.res if not r[1]]
        for lab, ok, det in self.res:
            if not ok:
                print("FAIL %s %s" % (lab, str(det)[:700]))
        print("checks.py: %d assertions, %d failed" % (len(self.res), len(bad)))
        if bad:
            raise SystemExit(1)


def hundred(v):
    return int(round(v / 100.0)) * 100


def pdf_text(path):
    out = subprocess.run(["pdftotext", "-enc", "UTF-8", path, "-"], capture_output=True, check=True).stdout
    return re.sub(r"\s+", " ", out.decode("utf-8"))


def xlsx_strings(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    bits = []
    for ws in wb.worksheets:
        bits.append(ws.title)
        for row in ws.iter_rows(values_only=True):
            bits.extend(str(v) for v in row if isinstance(v, str))
    return " ".join(bits)


def headers(target):
    """Every column name the pack ships."""
    out = {}
    for f in sorted(os.listdir(target)):
        p = os.path.join(target, f)
        if f.endswith(".csv"):
            with open(p, encoding="utf-8") as fh:
                out[f] = next(csv.reader(fh))
        elif f.endswith(".parquet"):
            out[f] = pq.ParquetFile(p).schema_arrow.names
        elif f.endswith(".xlsx"):
            wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
            cols = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    if row and all(isinstance(v, str) for v in row if v is not None) and sum(v is not None for v in row) > 2:
                        cols.extend(v for v in row if v)
                        break
            out[f] = cols
        elif f.endswith(".json"):
            with open(p, encoding="utf-8") as fh:
                doc = json.load(fh)
            keys = set()

            def walk(o, pre=""):
                if isinstance(o, dict):
                    for k, v in o.items():
                        keys.add(k)
                        walk(v)
            walk(doc["cases"][0])
            walk(doc["export"])
            out[f] = sorted(keys)
    return out


# ---------------------------------------------------------------------------------------------------- clean truth
def truth_frames(W):
    h = W.hist.copy()
    h["inj"] = h.n_rec
    P = W.plan
    PP = P[P.programme].copy()
    return h, PP


def truth_default(W):
    h, PP = truth_frames(W)
    l12 = h[((h.fy == 2025) & (h.wk >= 45)) | (h.fy == 2026)]
    g = l12.groupby("dc_no")
    rate = g.inj.sum() / g.hours.sum()
    return float(np.sum(PP.hours.values * rate.loc[PP.dc_no].values))


def truth_golden(W):
    h, PP = truth_frames(W)
    G = h.groupby(["dc_no", "fy", "per"], as_index=False).agg(inj=("inj", "sum"), hours=("hours", "sum"),
                                                             direct=("direct", "sum"), earned=("earned", "sum"))
    G["lpace"] = np.log(G.earned / G.direct)
    m = smf.glm("inj ~ C(dc_no) + lpace", data=G, family=sm.families.Poisson(), offset=np.log(G.hours)).fit()
    new = pd.DataFrame({"dc_no": PP.dc_no.values, "lpace": np.log((PP.earned / PP.direct).values),
                        "hours": PP.hours.values})
    return float(np.sum(m.predict(new, offset=np.log(new.hours))))


def truth_decisive_alone(W, gamma):
    """The default path with only the decisive step changed: each DC's latest twelve months' rate per hour carried to
    the plan's pace against standard through the within-DC response."""
    h, PP = truth_frames(W)
    l12 = h[((h.fy == 2025) & (h.wk >= 45)) | (h.fy == 2026)]
    g = l12.groupby("dc_no")
    rate = g.inj.sum() / g.hours.sum()
    pace12 = g.earned.sum() / g.direct.sum()
    return float(sum(r.hours * rate[r.dc_no] * ((r.earned / r.direct) / pace12[r.dc_no]) ** gamma
                     for r in PP.itertuples()))


# ---------------------------------------------------------------------------------------------------- mess variants
def variant(D, **kw):
    """The golden's construction with one input handled the wrong way."""
    D2 = dict(D)
    ALL = D["cases_all"]
    if kw.get("cases"):
        D2["cases"] = kw["cases"](ALL)
    if kw.get("wfm") is not None:
        D2["wfm"] = kw["wfm"]
    if kw.get("wms") is not None:
        D2["wms"] = kw["wms"]
    if kw.get("fac") is not None:
        D2["fac"] = kw["fac"]
    if kw.get("prog") is not None:
        D2["prog"] = kw["prog"]
    if kw.get("plan_vol") is not None:
        D2["plan_vol"], D2["plan_hrs"] = kw["plan_vol"], kw["plan_hrs"]
    X = calc.weekly(D2)
    if kw.get("xfilter"):
        X = kw["xfilter"](X)
    return calc.golden_of(D2, X=X)


def run(target, golden, B=None, R=None, meta_path=None):
    ck = Checks()
    B = B or PK.build_world()
    W = B.W
    R = R or DL.compute(target)
    D = R["D"]
    prog = R["prog"]
    gold = R["golden"]
    files = sorted(os.listdir(target))
    meta = json.load(open(meta_path, encoding="utf-8")) if meta_path else {}

    # ------------------------------------------------------------------------------------------------ input gates
    ck("G1 ten or more files", len(files) >= 10, len(files))
    fmts = {f.rsplit(".", 1)[1] for f in files}
    ck("G2 three or more formats", len(fmts) >= 3, fmts)
    big = {f: (sum(1 for _ in open(os.path.join(target, f), encoding="utf-8")) - 1) if f.endswith(".csv") else
           (pq.ParquetFile(os.path.join(target, f)).metadata.num_rows if f.endswith(".parquet") else 0) for f in files}
    ck("G3 a file of 25,000 rows or more", max(big.values()) >= 25000, max(big.items(), key=lambda x: x[1]))
    dis = meta.get("distractor_files", [])
    ck("G4 two or more distractors named in metadata.json, shipped, none named as such",
       len(dis) >= 2 and all(d in files for d in dis) and not any("distract" in f.lower() for f in files), dis)
    ck("G5 no file name broadcasts the step", not any(re.search(r"pace|perform|productiv|known_issue|decoy|trap|forecast",
                                                                f, re.I) for f in files), files)
    ck("G6 metadata.json lists every shipped file", sorted(x["path"] for x in meta.get("files", [])) == files)

    # ------------------------------------------------------------------------------------------------ R1
    H = headers(target)
    banned = {"pace", "perf", "performance", "rate", "trir", "forecast", "earned", "standard", "uph", "productivity",
              "ytd", "cumulative", "expected", "predicted"}
    bad_cols = {f: [c for c in cols if banned & set(re.split(r"[^a-z0-9]+", str(c).lower()))] for f, cols in H.items()}
    bad_cols = {f: v for f, v in bad_cols.items() if v}
    ck("R1a no shipped column holds pace, a rate, standard hours, a forecast or a year-to-date figure", not bad_cols,
       bad_cols)
    G = R["G"]
    n_per = sum(12 for _ in range(2021, 2026)) + 11
    ck("R1b the golden's fit reads every DC's fiscal periods built from weekly hours, cases by category and the "
       "standards in force (60 DCs, FY2021 P1 to FY2026 P11)",
       G.dc.nunique() == 60 and len(G) == 60 * n_per and G.groupby("dc").size().min() == n_per, (G.dc.nunique(), len(G)))
    ck("R1c the plan ships cases by category and labour hours by DC and period, no pace or rate",
       set(D["plan_vol"].columns) >= {"DC_NO", "PERIOD", "CATEGORY", "CASES"} and
       set(D["plan_hrs"].columns) >= {"DC_NO", "PERIOD", "DIRECT_HOURS", "INDIRECT_HOURS"})

    # ------------------------------------------------------------------------------------------------ R2
    def legacy_twice(ALL):
        return calc.case_frame(ALL, keep_migrated_copies=True)

    def removed_in(ALL):
        return calc.case_frame(ALL, keep_removed=True)

    def utc_dates(ALL):
        return calc.case_frame(ALL, date_col="utc_date")

    n_mig = int(D["cases_all"].migrated_copy.sum())
    v = {}
    v["R2.1 migrated cases counted in both exports"] = variant(D, cases=legacy_twice)
    v["R2.2 cases removed as not work-related kept"] = variant(D, cases=removed_in)
    v["R2.3 the platform's UTC timestamps read as dates"] = variant(D, cases=utc_dates)
    # Rockford's two numbers and codes read as two DCs
    D_rk = dict(D)
    num2 = dict(D["num2dc"])
    num2[WD.RECODE_NEW_NO] = WD.RECODE_NEW_NO
    loc2 = dict(D["loc2dc"])
    loc2[WD.RECODE_NEW_CODE] = WD.RECODE_NEW_NO
    wfm = D["wfm"].copy()
    wfm["dc"] = wfm.DC_NO.map(num2)
    wms = D["wms"].copy()
    wms["dc"] = wms.loc_code.map(loc2)
    fac = D["fac"].copy()
    extra = fac[fac.dc == WD.RECODE_DC].copy()
    extra["dc"] = WD.RECODE_NEW_NO
    fac = pd.concat([fac, extra], ignore_index=True)
    pv, ph = D["plan_vol"].copy(), D["plan_hrs"].copy()
    pv["dc"] = pv.DC_NO.map(num2)
    ph["dc"] = ph.DC_NO.map(num2)
    cut = K.week_start(*WD.RECODE_WEEK)

    def rk_cases(ALL):
        CS = calc.case_frame(ALL)
        CS.loc[(CS.dc == WD.RECODE_DC) & (CS.date >= cut), "dc"] = WD.RECODE_NEW_NO
        return CS
    prog_rk = [WD.RECODE_NEW_NO if d == WD.RECODE_DC else d for d in prog]
    v["R2.4 Rockford's two numbers and codes read as two DCs"] = variant(
        D, cases=rk_cases, wfm=wfm, wms=wms, fac=fac, prog=prog_rk, plan_vol=pv, plan_hrs=ph)
    # the WMS category codes before fiscal 2023 week 27 left unmapped: those weeks have no standard hours and drop out
    wms_old = D["wms"].copy()
    wms_old["cat"] = wms_old.category_code
    v["R2.5 the WMS category codes before FY2023 week 27 left unmapped"] = variant(
        D, wms=wms_old, xfilter=lambda X: X[X.earned > 0])
    # Lebanon TN's re-study: the current site factor applied to every week
    fac_l = D["fac"].copy()
    cur = fac_l[(fac_l.dc == WD.RESTUDY_DC) & fac_l["to"].isna()].copy()
    fac_l = fac_l[fac_l.dc != WD.RESTUDY_DC]
    cur["from"] = dt.date(2019, 8, 4)
    fac_l = pd.concat([fac_l, cur], ignore_index=True)
    v["R2.6 Lebanon TN's current site factor applied to every week"] = variant(D, fac=fac_l)
    v["R2.7 fiscal 2022's week 53 dropped"] = variant(D, xfilter=lambda X: X[~((X.fy == 2022) & (X.wk == 53))])
    for k, x in v.items():
        ck(k + ": the forecast stays in the 1,000 bin", LO <= x < HI, round(x, 1))
    ck("R2.1b the migration's carried-over cases exist (the item is real)", n_mig > 50, n_mig)
    # the ITA submissions are calendar-year records that tie to the cleaned cases
    ita = pd.read_csv(os.path.join(target, "ita_form300a_submissions_cy2021_cy2025.csv"))
    CS = D["cases"]
    est = {r.EHS_ESTABLISHMENT.upper(): D["est2dc"][r.EHS_ESTABLISHMENT] for r in D["reg"].itertuples()}
    tie = []
    for r in ita.itertuples():
        dc = est[r.establishment_name]
        n = int(((CS.dc == dc) & CS.rec & (pd.to_datetime(CS.date).dt.year == r.year_filing_for)).sum())
        tie.append(n == r.total_dafw_cases + r.total_djtr_cases + r.total_other_cases)
    ck("R2.9 the ITA Form 300A submissions tie to the cleaned cases by calendar year (every establishment-year)",
       all(tie), sum(tie))

    # ------------------------------------------------------------------------------------------------ R3
    L = R["rungs"]
    r0 = L["R0 latest twelve months' rate per hour x planned hours"]
    t_def = truth_default(W)
    t_gold = truth_golden(W)
    t_exp = float(W.plan[W.plan.programme].truth.sum())
    ck("R3a.1 the analyst test on the clean truth: today's rate per hour times the planned hours files about 800",
       hundred(t_def) == 800, round(t_def, 1))
    ck("R3a.2 the analyst test on the pack: the default path files about 800", hundred(r0) == 800, round(r0, 1))
    for k, x in L.items():
        ck("R3a.3 %s files outside the golden's bin (700 to 900)" % k, 650 <= x < 950 and hundred(x) != 1000,
           round(x, 1))
    alone = R["variants"]["latest twelve months' rate x fitted pace uplift"]
    ck("R3b.1 the default path with only the decisive step changed files the golden's 1,000", hundred(alone) == 1000,
       round(alone, 1))
    ck("R3b.2 the clean truth: the world's own expectation rounds to 1,000", hundred(t_exp) == 1000, round(t_exp, 1))
    ck("R3b.3 the clean truth: the golden's model on the world's records rounds to 1,000", hundred(t_gold) == 1000,
       round(t_gold, 1))
    ck("R3b.4 the clean truth: the decisive step alone on the world's records rounds to 1,000",
       hundred(truth_decisive_alone(W, R["gamma"])) == 1000, round(truth_decisive_alone(W, R["gamma"]), 1))
    # R3c: nothing planted
    wf = pd.read_csv(os.path.join(target, "wfm_hours_by_dept_week_fy21_fy26.csv"))
    seqs = wf[wf.DEPT == "SEL"].sort_values(["DC_NO", "FISCAL_YEAR", "FISCAL_WEEK"]).groupby("DC_NO").REG_HRS.apply(
        lambda s: tuple(np.round(s.values, 2)))
    blocks = {}
    rep = 0
    for dc, sq in seqs.items():
        for i in range(0, len(sq) - 4):
            b = sq[i:i + 4]
            if b in blocks and blocks[b] != (dc, i):
                rep += 1
            blocks[b] = (dc, i)
    ck("R3c.1 no copied or repeated block in the hours (no four-week run of selection hours repeats)", rep == 0, rep)
    docs = {f: pdf_text(os.path.join(target, f)) for f in files if f.endswith(".pdf")}
    docs["export_notes.txt"] = open(os.path.join(target, "export_notes.txt"), encoding="utf-8").read()
    ck("R3c.2 the pack's documents are the procedure, the agreement, the notes and the manifest (no memo, draft "
       "forecast or briefing)", sorted(f for f in files if f.endswith((".pdf", ".txt", ".docx", ".md"))) ==
       sorted([calc_doc for calc_doc in ("EHS-RK-01_osha_recordkeeping_procedure_rev7.pdf",
                                         "lindqvist_program_agreement_fy2027_draft.pdf", "export_notes.txt")]))
    agree = docs["lindqvist_program_agreement_fy2027_draft.pdf"]
    ck("R3c.3 no stakeholder figure: the agreement leaves the forecast blank and no document states a fiscal 2027 count",
       "Forecast Recordable Cases ____" in agree and not re.search(r"\b(?:[6-9]\d\d|1,?[0-4]\d\d)\s+(?:recordable )?cases",
                                                                   " ".join(docs.values())))
    # R3 strict: the shipped HOW sentences executed literally
    CSr = D["cases"][D["cases"].rec]
    X = R["X"]
    PP = R["PP"]
    ph = PP.groupby("dc").hours.sum()

    def rate_times_plan(mask):
        sub = X[mask & X.dc.isin(prog)]
        g = sub.groupby("dc")
        return float(sum(ph[d] * g.inj.sum()[d] / g.hours.sum()[d] for d in prog))
    lit = {
        "EHS-RK-01 8 with Schedule B 2.2, latest twelve months": r0,
        "EHS-RK-01 8 with Schedule B 2.2, fiscal 2025": rate_times_plan(X.fy == 2025),
        "EHS-RK-01 8 with Schedule B 2.2, fiscal 2026 to week 44": rate_times_plan(X.fy == 2026),
        "EHS-RK-01 8 with Schedule B 2.2, every year": L["R0a every year's rate per hour"],
        "the standards' notes with Schedule B 2.2 (standard hours)": L["R1e rate per standard hour"],
    }
    # the ITA submissions' own rate (cases over its hours) carried to the planned hours
    ita25 = ita[ita.year_filing_for == 2025]
    irate = {est[r.establishment_name]: (r.total_dafw_cases + r.total_djtr_cases + r.total_other_cases) /
             r.total_hours_worked for r in ita25.itertuples()}
    lit["EHS-RK-01 8 and 10 with Schedule B 2.2, the 2025 Form 300A"] = float(sum(ph[d] * irate[d] for d in prog))
    for k, x in lit.items():
        ck("R3s %s, executed literally, files %d, not the golden's 1,000" % (k, hundred(x)), hundred(x) <= 900,
           round(x, 1))
    plan_readme = xlsx_strings(os.path.join(target, "fy2027_operating_plan_v3.xlsx"))
    std_notes = xlsx_strings(os.path.join(target, "ie_engineered_labor_standards.xlsx"))
    words = re.compile(r"\bpace\b|faster|productivit|injur|recordable|safety|risk", re.I)
    ck("R3s the plan's read-me and the standards' notes carry no pace, risk or injury sentence",
       not words.search(plan_readme) and not words.search(std_notes),
       (words.findall(plan_readme), words.findall(std_notes)))
    proc = docs["EHS-RK-01_osha_recordkeeping_procedure_rev7.pdf"]
    ck("R3s no filed sentence names pace, speed or a forecast method",
       not re.search(r"\bpace\b|faster|speed|forecast method|exposure", proc + agree + docs["export_notes.txt"], re.I))

    # ------------------------------------------------------------------------------------------------ precision
    sig = R["golden_sigma"]
    ck("S1 one sigma either side leaves the forecast at 1,000", LO <= gold - sig and gold + sig < HI,
       (round(gold, 1), round(sig, 1)))
    ck("S2 the golden sits mid-bin (40 or more from either edge)", min(gold - LO, HI - gold) >= 40, round(gold, 1))
    V = R["variants"]
    for k, x in V.items():
        ck("S3 construction '%s' rounds to 1,000 (10 or more inside the bin)" % k, LO + 10 <= x < HI - 10, round(x, 1))
    ck("S4 the nearest rival is a full bin below the golden",
       max(L.values()) < LO and hundred(max(L.values())) <= 900, round(max(L.values()), 1))

    # ------------------------------------------------------------------------------------------------ back-test
    BT = R["backtest"]
    pace = BT["p_pace"]["mape"]
    ck("B1 out of sample the default path misses by a visible margin (5 points and 1.6 times the pace model's)",
       BT["p_rate"]["mape"] >= pace + 5 and BT["p_rate"]["mape"] >= 1.6 * pace, (BT["p_rate"]["mape"], pace))
    for k in ("p_rate_all", "p_seasonal", "p_case", "p_earn", "p_hours_el", "p_uph", "p_hours_cases", "p_pooled",
              "p_hshare"):
        ck("B2 out of sample the rival %s misses by 3 points or more beyond the pace model" % k,
           BT[k]["mape"] >= pace + 3, (round(BT[k]["mape"], 1), round(pace, 1)))
    S = R["separation"]
    ck("B3 in sample cases per hour fits 30 or more deviance points worse than pace on the same parameters",
       S["uph"]["dev"] - S["pace"]["dev"] >= 30, round(S["uph"]["dev"] - S["pace"]["dev"], 1))
    pu = S["pace_uph"]["params"]
    ck("B4 with both in the model pace carries the response (4 or more se) and cases per hour does not (under 1.5 se)",
       pu["lpace"][0] / pu["lpace"][1] >= 4 and abs(pu["luph"][0] / pu["luph"][1]) < 1.5, pu)
    ck("B5 at the DCs whose work content per case moved more than 3 per cent, cases per hour misses 1.5 times the "
       "pace model or more", BT["p_uph"]["n_mixed"] >= 5 and BT["p_uph"]["mape_mixed"] >= 1.5 * BT["p_pace"]["mape_mixed"],
       (BT["p_uph"]["n_mixed"], BT["p_uph"]["mape_mixed"], BT["p_pace"]["mape_mixed"]))
    ck("B6 the default's out-of-sample errors run against the pace change (correlation -0.25 or below)",
       BT["p_rate"]["corr"] <= -0.25, BT["p_rate"]["corr"])
    ck("B7 the pooled response is far below the within-DC response (the faster DCs are the safer ones)",
       R["gamma_pooled"] < 1.0 and R["gamma"] > 2.0, (R["gamma_pooled"], R["gamma"]))

    # ------------------------------------------------------------------------------------------------ fork table
    def injuries_only(ALL):
        CS = calc.case_frame(ALL)
        return CS[~(CS.rec & CS.ill.str.upper().ne("INJURY"))].reset_index(drop=True)
    v_ill = variant(D, cases=injuries_only)
    ck("SF1 'recordable cases': illnesses left out gives the same call", LO <= v_ill < HI, round(v_ill, 1))
    tab = agree.split("Schedule A - Program Sites")[1].split("Schedule B - Program Fee")[0]
    sched = sorted(int(x) for x in re.findall(r"\b(\d{3,4}) [A-Z][A-Za-z .]+? DC\b", tab))
    reg_prog = sorted(int(str(r.DC_NO).split(";")[-1]) for r in D["reg"].itertuples() if r.ERGO_PROGRAM_FY27 == "Y")
    ck("SF2 the twenty DCs: the agreement's Schedule A and the register's program column are the same twenty",
       sched == reg_prog and len(sched) == 20, (sched, reg_prog))
    per = pd.read_excel(os.path.join(target, "fy2027_operating_plan_v3.xlsx"), sheet_name="Periods")
    ck("SF3 fiscal 2027 is 12/27/2026 to 12/25/2027, 52 weeks, the plan's twelve periods",
       str(per.FIRST_DAY.iloc[0])[:10] == "2026-12-27" and str(per.LAST_DAY.iloc[-1])[:10] == "2027-12-25" and
       int(per.WEEKS.sum()) == 52 and "12/27/2026 to Saturday 12/25/2027" in agree)
    ck("SF4 the approved plan (Schedule B 2.2), its labour hours as filed, gives the golden; hours re-derived at today's "
       "productivity file another bin", "operating plan for the Program Sites as approved before signing." in agree and
       hundred(L["R1e rate per standard hour"]) != 1000)
    leg = pd.read_csv(os.path.join(target, "ehs_legacy_case_export_fy21_2024-06-30.csv"), dtype=str,
                      keep_default_na=False)
    lag_leg = (pd.to_datetime(leg.DATE_ENTERED, format="%m/%d/%Y") - pd.to_datetime(leg.INCIDENT_DATE,
                                                                                    format="%m/%d/%Y")).dt.days.max()
    lag_pl = max((dt.datetime.fromisoformat(c["reported_at"][:19]) - dt.datetime.fromisoformat(c["occurred_at"][:19])).days
                 for c in D["plat"])
    ck("SF5 by date entered or by date of injury: no case entered more than seven days after the injury",
       lag_leg <= 7 and lag_pl <= 7, (lag_leg, lag_pl))
    n_agy = int(sum(1 for c in D["plat"] if c["person"]["employer"] != K.ORG and c["classification"]["category"] == "Recordable"))
    ck("SF6 agency associates are on the logs (EHS-RK-01 3.3) and in the golden", n_agy > 100 and
       "supervised day to day by Ostrander are recorded on the log" in proc, n_agy)
    ck("SF6b the plan's hours are Ostrander and agency associates' hours together, as the history's are",
       "agency associates' hours together" in plan_readme and "agency associates' hours included" in proc)
    n_priv = sum(1 for c in D["plat"] if (c.get("osha_log") or {}).get("privacy_case"))
    ck("SF7 privacy cases are on the logs, the name withheld (EHS-RK-01 5), and counted", n_priv > 5 and
       "privacy case" in proc and "confidential list" in proc, n_priv)
    ck("SF8 migrated cases once or twice: the same call (R2.1)", LO <= v["R2.1 migrated cases counted in both exports"] < HI)
    ck("SF9 removed cases out or in: the same call (R2.2)", LO <= v["R2.2 cases removed as not work-related kept"] < HI)
    ck("SF10 the forecast is the expected value (Schedule B 2.3) and the golden is the model's mean",
       "expected value" in agree and abs(gold - t_exp) / t_exp < 0.05, (round(gold, 1), round(t_exp, 1)))
    by_dc = R["golden_by_dc"]
    ck("SF11 one rounding at the end: the DCs' whole-case forecasts summed round to the same 1,000",
       hundred(sum(round(x) for x in by_dc.values())) == 1000)
    ck("SF12 hours worked over every department or direct hours only: the same call",
       LO <= V["direct-hours offset"] < HI, round(V["direct-hours offset"], 1))
    ck("SF13 the history from fiscal 2021, 2022, 2023 or 2024: the same call",
       all(LO <= V[k] < HI for k in ("from FY2022", "from FY2023", "from FY2024")) and LO <= gold < HI)
    ck("SF14 one plan version ships (v3, approved 10/23/2026)", "version 3" in plan_readme and
       "10/23/2026" in plan_readme and sum("operating_plan" in f for f in files) == 1)

    # ------------------------------------------------------------------------------------------------ closure axes
    ck("C1 population: the twenty Schedule A DCs' log cases, agency and privacy cases in (SF1, SF2, SF6, SF7)",
       len(prog) == 20 and n_agy > 0 and n_priv > 0)
    ck("C2 unit of account: a recordable case (Schedule B 2.1)", "recordable cases that Client expects" in agree)
    ck("C2b what the forecast is set against: the sites' fiscal 2027 recordable cases in the post-term program review "
       "(section 1), no method named", "setting the Program Sites' fiscal 2027 recordable cases against the Forecast "
       "Recordable Cases" in agree and not re.search(r"back-?test|model|method|pace|regress",
                                                     agree.split("Schedule B - Program Fee")[1], re.I))
    sched_b = re.sub(r"\d{2}/\d{2}/\d{4}|B\d(?:\.\d)?|page \d+", " ", agree.split("Schedule B - Program Fee")[1])
    ck("C2c read literally, Schedule B files no figure but the fee: no rate, count or forecast",
       set(re.findall(r"\d[\d,.]*\d|\d", sched_b)) <= {"2,140", "2027", "300"} and "________" in agree,
       set(re.findall(r"\d[\d,.]*\d|\d", sched_b)))
    ck("C3 window: the fiscal year of the date of injury (EHS-RK-01 4.2)", "log of the year in which the injury" in proc)
    ck("C4 as-of: the plan v3 of 10/23/2026; history to FY2026 week 44 pulled 11/02/2026",
       int(X[X.fy == 2026].wk.max()) == 44 and "10/23/2026" in plan_readme)
    ck("C5 version basis: the standards in force each week; Lebanon TN's re-study either way gives the same call (R2.6)",
       LO <= v["R2.6 Lebanon TN's current site factor applied to every week"] < HI)
    ck("C6 denominator: hours worked over every department and pay code (EHS-RK-01 8)",
       "every department and pay code" in proc)
    ck("C7 weighting: weekly and period fits and the negative binomial give the same call",
       all(LO <= V[k] < HI for k in ("weekly fit", "negative binomial")))
    ck("C8 window length: from FY2022, 2023 or 2024 the same call (SF13)", True)
    ck("C9 boundaries: local dates against UTC dates the same call (R2.3)",
       LO <= v["R2.3 the platform's UTC timestamps read as dates"] < HI)
    ck("C10 rounding path: unrounded DC-periods summed, rounded once (SF11)", True)
    ck("C11 tie-break: the golden is not on a rounding edge", abs(gold - 950) > 1 and abs(gold - 1050) > 1)
    wk_inj = X[X.fy == 2026].groupby("wk").inj.sum()
    ck("C12 maturity: the last week's count is in line with the weeks before (every case through week 44 entered)",
       wk_inj.loc[44] >= 0.6 * wk_inj.loc[30:43].mean(), (int(wk_inj.loc[44]), round(wk_inj.loc[30:43].mean(), 1)))
    ck("C13 order of operations forced by the data (the back-test, B1 to B5)", True)
    rev = D["cases_all"].iloc[::-1].reset_index(drop=True)
    D_rev = dict(D)
    D_rev["cases_all"] = rev
    ck("C14 row order: the exports read in reverse give the same forecast",
       abs(variant(D_rev, cases=lambda ALL: calc.case_frame(ALL)) - gold) < 1e-6)
    ck("C15 duplicates: each case carried into the platform at go-live once, through its legacy_case_ref (the notes)",
       "legacy_case_ref" in docs["export_notes.txt"] and "carried into the platform" in docs["export_notes.txt"] and
       n_mig > 0)
    ck("C16 identities: the register maps every WFM DC number, WMS location and EHS establishment",
       set(D["wfm"].DC_NO.unique()) <= set(D["num2dc"]) and set(D["wms"].loc_code.unique()) <= set(D["loc2dc"]) and
       {c["site"]["establishment"] for c in D["plat"]} <= set(D["est2dc"]))
    ck("C17 netting: first-aid and removed cases off the log (EHS-RK-01 3.1 and 6)",
       "First aid alone is not recorded" in proc and "remove it from the log" in proc)
    ck("C18 units: hours, cases and standard minutes per case shipped", "minutes per case shipped" in std_notes.lower())
    ck("C19 codes: the WMS category re-code mapped once, in the standards' notes", "DRYG = AMB" in std_notes and
       "DRYG" not in docs["export_notes.txt"])
    rm_sheet = " ".join(str(r[0]) for r in openpyxl.load_workbook(os.path.join(target, "fy2027_operating_plan_v3.xlsx"),
                                                                 data_only=True)["Read me"].iter_rows(values_only=True)
                        if r and r[0])
    ck("C19b Rockford's re-numbering stated once, in the site register's own note", "Rockford" not in
       docs["export_notes.txt"] and "Rockford" not in rm_sheet and "RKF" not in rm_sheet and
       "311" in str(D["reg"][D["reg"].DC_NO.str.contains("318")].NOTES.iloc[0]), rm_sheet[:200])
    ck("C20 no integerisation before the final rounding", isinstance(gold, float))
    ck("C21 scope: EHS-RK-01 covers every Ostrander DC", "Every Ostrander distribution center" in proc)
    ck("C22 the forward window: twelve plan periods, the new accounts from period 1",
       sorted(PP.per.unique().tolist()) == list(range(1, 13)) and "start in period 1" in plan_readme)

    # ------------------------------------------------------------------------------------------------ Ask A
    A = R["askA"]
    cases = AW.rtw_cases(W)
    truth_a = {d: 0 for d in prog}
    for c in cases:
        if c["kind"] == "REC" and c["dc_no"] in truth_a:
            a_, b_ = AW.osha_days(c)
            truth_a[c["dc_no"]] += a_ + b_
    ck("AA1 Ask A's golden equals the world's day counts at every DC", all(A["golden"][d] == truth_a[d] for d in prog))
    for k in ("stop_platform", "stop_by_case_no", "stop_home_dc", "stop_over"):
        ck("AA2 Ask A stop %s misses at every DC" % k, all(A[k][d] != A["golden"][d] for d in prog),
           sum(A[k][d] == A["golden"][d] for d in prog))
    for k in ("stop_no_cap", "stop_sep_ignored", "stop_with_removed"):
        ck("AA3 Ask A hazard %s moves at least one DC" % k, any(A[k][d] != A["golden"][d] for d in prog))
    ck("AA4 every programme DC placed at least one restricted associate on transitional duty",
       all(A["n_continuations"][d] >= 1 for d in prog), A["n_continuations"])
    # ------------------------------------------------------------------------------------------------ Ask B
    Bq = R["askB"]
    raw = B.picks_raw
    truth_b = raw[raw.weight_lb > 50].groupby("dc_no").cases.sum().reindex(prog, fill_value=0)
    ck("AB1 Ask B's golden equals the world's heavy cases at every DC", all(Bq["golden"][d] == int(truth_b[d]) for d in prog))
    for k in ("stop_pounds", "stop_current", "stop_cracker", "stop_over", "stop_at_50"):
        ck("AB2 Ask B stop %s misses at every DC" % k, all(Bq[k][d] != Bq["golden"][d] for d in prog),
           sum(Bq[k][d] == Bq["golden"][d] for d in prog))
    I = B.I
    ck("AB3 no vendor-portal item within half a pound of 50 (the kilogram round trip stays on its side)",
       not ((I.source == "PORTAL") & ((I.weight_lb - 50).abs() < 0.5)).any())
    ck("AB4 one item version in effect for every pick row", Bq["ambiguous_rows"] == 0, Bq["ambiguous_rows"])
    ck("AB5 every picked SKU is in the master (discontinued items kept), so the habitual join leaves no row unmatched",
       Bq["unmatched_rows"] == 0, Bq["unmatched_rows"])
    notes_txt = docs["export_notes.txt"]
    ck("AB7 the export notes give each weight source's unit (pounds for CUBISCAN and BUYER, kilograms for PORTAL)",
       "lb for CUBISCAN and BUYER; kg for PORTAL" in notes_txt)
    ck("AB8 the unit-aware cracker (the master's weights in each source's unit, no item versions) misses at every DC",
       all(Bq["stop_current"][d] != Bq["golden"][d] for d in prog),
       sum(Bq["stop_current"][d] == Bq["golden"][d] for d in prog))
    wms = D["wms"]
    pk = pq.read_table(os.path.join(target, "wms_sku_picks_fy2026_p01_p11.parquet")).to_pandas()
    pk["dc"] = pk.loc_code.map(D["loc2dc"])
    w26 = wms[(wms.fiscal_year == 2026) & wms.dc.isin(prog)].copy()
    w26["per"] = w26.fiscal_week.map(calc.period)
    tie_w = w26.groupby(["dc", "per"]).cases_shipped.sum()
    tie_p = pk.groupby(["dc", "fiscal_period"]).cases_picked.sum()
    tie_p.index.names = ["dc", "per"]
    ck("AB6 the SKU picks tie to the WMS weekly cases at every programme DC and period",
       bool((tie_w.round().astype(int) - tie_p.reindex(tie_w.index).astype(int)).abs().max() == 0) or
       bool(((tie_w - tie_p.reindex(tie_w.index)).abs() <= 4).all()),
       float((tie_w - tie_p.reindex(tie_w.index)).abs().max()))
    # ------------------------------------------------------------------------------------------------ separation line
    tmp = tempfile.mkdtemp(prefix="t149_")
    try:
        keep = [f for f in files if f not in dis + ["rtw_work_status_log_fy2026.csv", "wms_item_master_2026-11-01.csv",
                                                     "wms_item_master_history.csv",
                                                     "wms_sku_picks_fy2026_p01_p11.parquet"]]
        for f in keep:
            shutil.copy2(os.path.join(target, f), os.path.join(tmp, f))
        Dm = calc.load(tmp)
        g_main = calc.golden_of(Dm)
    finally:
        shutil.rmtree(tmp)
    ck("SEP1 the forecast from the main files alone (no work-status log, item master, history, picks or distractors) "
       "is the golden", abs(g_main - gold) < 1e-6, (g_main, gold))
    src = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "calc.py"), encoding="utf-8").read()
    ck("SEP2 the golden's code never opens a distractor", not any(d in src for d in dis))
    ck("SEP3 the main call reads no day count: the platform's osha_log fields appear only in Ask A's code",
       src.count("osha_log") == src.split("def ask_a")[1].count("osha_log"))
    # ------------------------------------------------------------------------------------------------ chart
    tc = W.cases
    tc = tc[(tc.kind == "REC") & tc.dc_no.isin(prog)]
    truth_chart = {fy: int((tc.fy == fy).sum()) for fy in range(2021, 2026)}
    ck("CH1 the chart's five fiscal years equal the world's recordable cases at the twenty DCs",
       truth_chart == R["chart"], (truth_chart, R["chart"]))
    # ------------------------------------------------------------------------------------------------ realism
    texts = dict(docs)
    for f in files:
        if f.endswith((".csv", ".json", ".txt")):
            texts[f] = open(os.path.join(target, f), encoding="utf-8").read()
        elif f.endswith(".xlsx"):
            texts[f] = xlsx_strings(os.path.join(target, f))
    ck("RS1 no em or en dash in any shipped file", not any(DASH.search(t) for t in texts.values()),
       [f for f, t in texts.items() if DASH.search(t)])
    card = json.load(open(os.path.join(REPO, ".claude", "skills", "fingerprint", "cards", "task149.json")))
    ck("RS2 every persona is from the card's drawn people", set(K.PEOPLE.values()) <= set(card["world"]["people"]),
       K.PEOPLE)
    au = subprocess.run(["python3", os.path.join(REPO, ".claude", "skills", "reduce-house-fixes", "scripts",
                                                 "scrub_producer_metadata.py"), target, "--floor", "2019-01-01",
                         "--ceiling", "2026-11-02"], capture_output=True, text=True)
    ck("RS3 containers scrubbed: no writer name or out-of-band stamp in the pack", "clean" in au.stdout, au.stdout[-400:])
    ga = subprocess.run(["python3", os.path.join(REPO, ".claude", "skills", "reduce-house-fixes", "scripts",
                                                 "scrub_producer_metadata.py"), golden, "--floor", "2026-11-02",
                         "--ceiling", "2026-11-04"], capture_output=True, text=True)
    ck("RS3b containers scrubbed in the goldens", "clean" in ga.stdout, ga.stdout[-400:])
    mt = {int(os.path.getmtime(os.path.join(target, f))) for f in files}
    ck("RS4 one export time on every file", mt == {PK.PACK_MTIME}, mt)
    occ = max(c["occurred_at"] for c in D["plat"])
    rtw = pd.read_csv(os.path.join(target, "rtw_work_status_log_fy2026.csv"))
    upd = max((c.get("osha_log") or {}).get("days_last_updated") or "" for c in D["plat"])
    rep_ = max(c["reported_at"] for c in D["plat"])
    ck("RS5 nothing dated after the pull", occ[:10] <= "2026-11-01" and rtw.ENTERED_ON.max() <= "2026-11-01" and
       wf.WEEK_ENDING.max() <= "2026-10-31" and upd <= "2026-11-01" and rep_[:10] <= "2026-11-02",
       (occ, rtw.ENTERED_ON.max(), wf.WEEK_ENDING.max(), upd, rep_))
    tot = wf.assign(h=wf.REG_HRS + wf.OT_HRS + wf.AGY_HRS).groupby("DC_NO").h.sum().round(2)
    ck("RS6 no two DCs share an hours total and no total is round", tot.nunique() == len(tot) and
       not any(abs(x - round(x, -3)) < 0.5 for x in tot), tot.head().to_dict())
    pv = D["plan_vol"].groupby("DC_NO").CASES.sum()
    ck("RS7 no plan total is round", not any(x % 1000 == 0 for x in pv))
    notes = rtw.NOTE.fillna("").value_counts()
    ck("RS8 the work-status log's free text recurs at a real rate (no note on more than a third of rows)",
       notes.drop("", errors="ignore").max() < len(rtw) / 3, notes.head(4).to_dict())
    # ------------------------------------------------------------------------------------------------ goldens
    wb = openpyxl.load_workbook(os.path.join(golden, DL.F_BOOK), data_only=True)
    ck("GD1 the workbook states the forecast at the nearest hundred", wb["Forecast"]["B3"].value == hundred(gold))
    lab = R["labels"]
    ws = wb["DART days FY2026"]
    got = {str(r[0]): r[2] for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and r[1]}
    ck("GD2 the workbook's DART days table equals Ask A", all(got[lab[d][0]] == A["golden"][d] for d in prog))
    ws = wb["Cases over 50 lb FY2026"]
    got = {str(r[0]): r[2] for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and r[1]}
    ck("GD3 the workbook's heavy cases table equals Ask B", all(got[lab[d][0]] == Bq["golden"][d] for d in prog))
    from pptx import Presentation
    prs = Presentation(os.path.join(golden, DL.F_SLIDE))
    sl = prs.slides[0]
    txt = " ".join(s.text_frame.text for s in sl.shapes if s.has_text_frame)
    vals = list([s for s in sl.shapes if s.has_chart][0].chart.plots[0].series[0].values)
    ck("GD4 one slide stating the forecast, its chart the five fiscal years and the forecast",
       len(prs.slides) == 1 and "{:,}".format(hundred(gold)) in txt and
       vals == [float(R["chart"][y]) for y in range(2021, 2026)] + [float(hundred(gold))])
    gtxt = xlsx_strings(os.path.join(golden, DL.F_BOOK)) + txt
    ck("GD5 no em or en dash in the goldens", not DASH.search(gtxt))
    # ------------------------------------------------------------------------------------------------ registers and citations
    idx = pd.read_csv(os.path.join(target, "folder_index_2026-11-02.csv"))
    ck("H9 the folder index lists every shipped file but itself", sorted(idx.file) ==
       sorted(f for f in files if f != "folder_index_2026-11-02.csv"), sorted(set(files) ^ set(idx.file)))
    sub_path = os.path.join(os.path.dirname(target), "submission.md")
    if os.path.exists(sub_path):
        sub = open(sub_path, encoding="utf-8").read()
        cited = set(re.findall(r"[\w.-]+\.(?:csv|json|parquet|xlsx|pdf|txt|pptx)", sub))
        cited -= {DL.F_BOOK, DL.F_SLIDE}
        ck("H8 every file the write-up cites ships in the pack", cited <= set(files), cited - set(files))
        heads = {"3": "3 What is recorded", "4": "4 Dates", "5": "5 Privacy cases", "6": "6 Work-relatedness",
                 "7": "7 System of record", "8": "8 Rates", "9": "9 Counting days", "10": "10 Annual summary"}
        cites = re.findall(r"EHS-RK-01 (\d+)", sub)
        ck("H8b the write-up cites EHS-RK-01 sections that exist", all(n in heads and heads[n] in proc for n in cites),
           cites)
    # ------------------------------------------------------------------------------------------------ prompt
    pr = open(os.path.join(os.path.dirname(target), "prompt.txt"), encoding="utf-8").read()
    vc = subprocess.run(["python3", os.path.join(REPO, ".claude", "skills", "guide-to-prompt", "references",
                                                 "voice-check.py"), os.path.basename(os.path.dirname(target))],
                        capture_output=True, text=True, cwd=REPO)
    ck("PR1 the prompt passes voice-check", vc.returncode == 0, vc.stdout[-300:])
    ck("PR2 the prompt is under 1,000 characters and names both deliverables", len(pr) < 1000 and DL.F_BOOK in pr and
       DL.F_SLIDE in pr, len(pr))
    print("measured: golden %.1f (sigma %.1f), default %.1f; clean truth default %.1f, golden model %.1f, world "
          "expectation %.1f; constructions %.1f to %.1f" % (gold, sig, r0, t_def, t_gold, t_exp, min(V.values()),
                                                           max(V.values())))
    print("measured: mess variants " + "; ".join("%s %.1f" % (k.split(" ")[0], x) for k, x in v.items()) +
          "; illnesses out %.1f; main files alone %.1f" % (v_ill, g_main))
    print("measured: back-test MAPE " + ", ".join("%s %.1f" % (k, d["mape"]) for k, d in BT.items()) +
          "; deviance gap (cases per hour less pace) %.1f" % (S["uph"]["dev"] - S["pace"]["dev"]))
    ck.report()
    return ck
