"""The golden deliverables, computed from the shipped pack only (calc.run): ergonomics_forecast_fy2027.xlsx (the
forecast and the two tables for the twenty DCs) and ergonomics_forecast_fy2027.pptx (one slide stating the forecast
with the five fiscal years' recordable cases charted beside it), and generator/figures.json, the figures the write-up
states, as it states them."""
from __future__ import annotations

import datetime as dt
import io
import json
import os
import re
import zipfile

import pandas as pd

import calc
import xlsx_util as XU

F_BOOK = "ergonomics_forecast_fy2027.xlsx"
F_SLIDE = "ergonomics_forecast_fy2027.pptx"
STAMP = dt.datetime(2026, 11, 4, 16, 38, 0)        # the analyst's save, two days after the pull
AUTHOR = "Jesse Delgado"


def n0(v):
    return "{:,}".format(int(round(v)))


def hundred(v):
    return int(round(v / 100.0)) * 100


def dc_labels(D):
    reg = D["reg"]
    out = {}
    for r in reg.itertuples():
        nums = [int(x) for x in str(r.DC_NO).split(";")]
        out[nums[0]] = (str(nums[-1]), r.EHS_ESTABLISHMENT)
    return out


def compute(target):
    R = calc.run(target)
    R["labels"] = dc_labels(R["D"])
    return R


def figures(R):
    lab = R["labels"]
    prog = R["prog"]
    A, B = R["askA"], R["askB"]
    bt = R["backtest"]
    S = R["separation"]
    rung = R["rungs"]
    var = R["variants"]
    f = {
        "call": {"forecast_recordable_cases": n0(hundred(R["golden"])),
                 "unrounded": "%.1f" % R["golden"], "dcs": "20", "fiscal_year": "2027"},
        "critical_components": {
            "pace_plan_vs_latest_twelve_months_per_cent": "%d" % round(100 * (R["pace_plan"] / R["pace_l12"] - 1)),
            "rate_per_hour_rise_per_cent": "%d" % round(100 * (R["rate_uplift"] - 1)),
            "response_to_log_pace": "%.1f" % R["gamma"],
        },
        "steps": {
            "latest_twelve_months_recordables": n0(R["l12_recordables_prog"]),
            "planned_hours_vs_latest_twelve_months_per_cent": "%d" % round(100 * (R["hours_ratio"] - 1)),
            "plan_cases_per_direct_hour_vs_latest_per_cent": "%d" % round(100 * (R["uph_ratio"] - 1)),
            "response_standard_error": "%.1f" % R["gamma_se"],
            "golden_standard_error": "%d" % round(R["golden_sigma"]),
            "backtest_pace_mape": "%d" % round(bt["p_pace"]["mape"]),
            "backtest_rate_mape": "%d" % round(bt["p_rate"]["mape"]),
            "backtest_uph_mape": "%d" % round(bt["p_uph"]["mape"]),
            "backtest_mixed_dcs": "%d" % bt["p_uph"]["n_mixed"],
            "backtest_pace_mape_mixed": "%d" % round(bt["p_pace"]["mape_mixed"]),
            "backtest_uph_mape_mixed": "%d" % round(bt["p_uph"]["mape_mixed"]),
            "deviance_gap_uph_vs_pace": "%d" % round(S["uph"]["dev"] - S["pace"]["dev"]),
            "constructions_low": n0(min(var.values())), "constructions_high": n0(max(var.values())),
        },
        "default_path": {"forecast": n0(hundred(rung["R0 latest twelve months' rate per hour x planned hours"])),
                         "unrounded": "%.1f" % rung["R0 latest twelve months' rate per hour x planned hours"]},
        "rivals": {k: n0(v) for k, v in rung.items()},
        "rivals_rounded": {k: n0(hundred(v)) for k, v in rung.items()},
        "fit": {"dcs": str(R["G"].dc.nunique()), "latest_twelve_months_from_week": "45",
                "latest_twelve_months_to_week": "44", "fy2022_weeks": "53"},
        "chart_fiscal_years": {str(k): n0(v) for k, v in R["chart"].items()},
        "ask_a_dart_days": [{"dc": lab[d][0], "site": lab[d][1], "days": n0(A["golden"][d])} for d in prog],
        "ask_a_total": n0(sum(A["golden"].values())),
        "ask_b_cases_over_50_lb": [{"dc": lab[d][0], "site": lab[d][1], "cases": n0(B["golden"][d])} for d in prog],
        "ask_b_total": n0(sum(B["golden"].values())),
        "ask_b_threshold_pounds": "50",
    }
    return f


# ---------------------------------------------------------------------------------------------------- workbook
def write_book(path, R):
    lab = R["labels"]
    prog = R["prog"]
    A, B = R["askA"], R["askB"]
    fc = hundred(R["golden"])
    wb = XU.Workbook()
    s = wb.sheet("Forecast")
    s.rows = [
        ["Fiscal 2027: about %s recordable cases forecast at the twenty ergonomics program DCs" % n0(fc)],
        ["Safety Analytics, 11/04/2026. Records through FY2026 week 44 (pulled 11/02/2026); FY2027 operating plan v3 "
         "(approved 10/23/2026)."],
        ["Forecast recordable cases, fiscal 2027 (12/27/2026 to 12/25/2027)", fc],
        ["Basis", "Expected cases with a date of injury in fiscal 2027 on the FY2027 operating plan v3 volumes and labor "
                  "hours, rounded to the nearest hundred (agreement Schedule B 2.1 to 2.3)"],
        ["Method", "Recordable rate per hour worked fitted on pace against the engineered labor standards (standard hours "
                   "earned on the WMS cases shipped over direct hours), each DC its own baseline, Poisson on DC fiscal "
                   "periods FY2021 to FY2026 wk 44, all 60 DCs; carried to the pace the plan's cases and hours set at "
                   "each program DC"],
        [],
        ["Recordable cases, twenty program DCs", "Cases"],
    ]
    for fy, v in sorted(R["chart"].items()):
        s.rows.append(["FY%d%s" % (fy, " (53 weeks)" if fy == 2022 else ""), v])
    s.rows.append(["FY2027 forecast", fc])
    s.widths = [60, 96]
    s.bold_rows = {1, 3, 7}
    s.header_row = 7
    s.number_formats = {2: "#,##0"}
    s.wrap_cols = {2}
    a = wb.sheet("DART days FY2026")
    a.rows = [["DC", "Site", "Days away, job transfer or restriction to 10/31/2026"]]
    for d in prog:
        a.rows.append([lab[d][0], lab[d][1], int(A["golden"][d])])
    a.rows.append(["Total", "", int(sum(A["golden"].values()))])
    a.rows.append([])
    a.rows.append(["Fiscal 2026 recordable cases with a date of injury 12/28/2025 to 10/31/2026 (EHS platform); calendar "
                   "days from the day after the injury to full duty, separation or 10/31/2026, 180 at most a case "
                   "(EHS-RK-01 9), from the return-to-work work-status log, transitional duty at a sister DC included"])
    a.widths = [8, 26, 50]
    a.bold_rows = {len(prog) + 2}
    a.number_formats = {3: "#,##0"}
    a.freeze = "A2"
    b = wb.sheet("Cases over 50 lb FY2026")
    b.rows = [["DC", "Site", "Cases picked over 50 lb, 12/28/2025 to 10/31/2026"]]
    for d in prog:
        b.rows.append([lab[d][0], lab[d][1], int(B["golden"][d])])
    b.rows.append(["Total", "", int(sum(B["golden"].values()))])
    b.rows.append([])
    b.rows.append(["WMS cases picked by SKU and fiscal period, FY2026 P1 to wk 44; case weight of the item version in "
                   "effect in the period picked (item master history), vendor-portal weights converted from kilograms; "
                   "cases over 50.0 lb"])
    b.widths = [8, 26, 50]
    b.bold_rows = {len(prog) + 2}
    b.number_formats = {3: "#,##0"}
    b.freeze = "A2"
    wb.save(path, title="FY2027 ergonomics program forecast", creator=AUTHOR, stamp=STAMP)


# ---------------------------------------------------------------------------------------------------- slide
def write_slide(path, R):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    from pptx.util import Emu, Inches, Pt

    fc = hundred(R["golden"])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    tb = sl.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(1.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "FY2027 forecast: about %s recordable cases at the twenty ergonomics program DCs" % n0(fc)
    p.font.size, p.font.bold = Pt(28), True
    p.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
    p2 = tf.add_paragraph()
    p2.text = ("Up from %s in FY2025: the operating plan moves the new accounts' heavier freight through hours worked "
               "about %d%% faster against the engineered standards than the latest twelve months, and each DC's "
               "recordable rate per hour rises with its pace" % (n0(R["chart"][2025]),
                                                                round(100 * (R["pace_plan"] / R["pace_l12"] - 1))))
    p2.font.size = Pt(15)
    p2.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    cd = CategoryChartData()
    cats = ["FY%d" % y for y in sorted(R["chart"])] + ["FY2027 forecast"]
    cd.categories = cats
    cd.add_series("Recordable cases", [R["chart"][y] for y in sorted(R["chart"])] + [fc])
    gf = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(2.0), Inches(11.7), Inches(4.6), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    plot = ch.plots[0]
    plot.gap_width = 70
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = "#,##0", False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(14)
    ser = plot.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = RGBColor(0x7F, 0x8C, 0x9E)
    pt = ser.points[len(cats) - 1]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = RGBColor(0xC0, 0x50, 0x2E)
    va = ch.value_axis
    va.has_major_gridlines = True
    va.maximum_scale, va.minimum_scale = 1200, 0
    va.tick_labels.font.size = Pt(12)
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = "#,##0", False
    va.has_title = True
    va.axis_title.text_frame.text = "Recordable cases"
    va.axis_title.text_frame.paragraphs[0].font.size = Pt(12)
    va.axis_title.text_frame.paragraphs[0].font.bold = False
    ch.category_axis.tick_labels.font.size = Pt(13)
    nb = sl.shapes.add_textbox(Inches(0.8), Inches(6.7), Inches(11.9), Inches(0.6))
    nt = nb.text_frame
    nt.word_wrap = True
    q = nt.paragraphs[0]
    q.text = ("Recordable cases on the twenty DCs' OSHA 300 Logs by fiscal year of the date of injury (FY2022 had 53 "
              "weeks). Forecast: expected cases with a date of injury in fiscal 2027 on the FY2027 operating plan v3, "
              "nearest hundred. Safety Analytics, 11/04/2026.")
    q.font.size = Pt(11)
    q.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    cp = prs.core_properties
    cp.author = AUTHOR
    cp.last_modified_by = AUTHOR
    cp.title = "FY2027 ergonomics program forecast"
    cp.created = STAMP
    cp.modified = STAMP
    cp.revision = 3
    cp.subject, cp.keywords, cp.comments, cp.category = "", "", "", ""
    buf = io.BytesIO()
    prs.save(buf)
    repack_pptx(buf.getvalue(), path)


def _props(body, name, author, app):
    iso = STAMP.strftime("%Y-%m-%dT%H:%M:%SZ")
    txt = body.decode("utf-8")
    if name.endswith("core.xml"):
        txt = re.sub(r"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)", r"\g<1>%s\g<2>" % iso, txt)
        txt = re.sub(r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", r"\g<1>%s\g<2>" % iso, txt)
        txt = re.sub(r"<dc:creator>[^<]*</dc:creator>", "<dc:creator>%s</dc:creator>" % author, txt)
        txt = re.sub(r"<dc:creator/>", "<dc:creator>%s</dc:creator>" % author, txt)
        if "<cp:lastModifiedBy" in txt:
            txt = re.sub(r"<cp:lastModifiedBy>[^<]*</cp:lastModifiedBy>",
                         "<cp:lastModifiedBy>%s</cp:lastModifiedBy>" % author, txt)
    else:
        txt = re.sub(r"<Application>[^<]*</Application>", "<Application>%s</Application>" % app, txt)
    return txt.encode("utf-8")


def _rezip(data, app, author):
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    for info in zin.infolist():
        body = zin.read(info.filename)
        if info.filename in ("docProps/core.xml", "docProps/app.xml"):
            body = _props(body, info.filename, author, app)
        elif info.filename.startswith("ppt/embeddings/") and info.filename.endswith(".xlsx"):
            body = _rezip(body, "Microsoft Excel", author)
        zi = zipfile.ZipInfo(info.filename, date_time=(STAMP.year, STAMP.month, STAMP.day, STAMP.hour, STAMP.minute, 0))
        zi.compress_type = zipfile.ZIP_DEFLATED
        zi.external_attr = 0o600 << 16
        zout.writestr(zi, body)
    zout.close()
    return out.getvalue()


def repack_pptx(data, path):
    with open(path, "wb") as f:
        f.write(_rezip(data, "Microsoft Office PowerPoint", AUTHOR))


def main(target, golden, figures_path, R=None):
    R = R or compute(target)
    os.makedirs(golden, exist_ok=True)
    write_book(os.path.join(golden, F_BOOK), R)
    write_slide(os.path.join(golden, F_SLIDE), R)
    F = figures(R)
    with open(figures_path, "w", encoding="utf-8") as fh:
        json.dump(F, fh, indent=1, ensure_ascii=False)
        fh.write("\n")
    return R
