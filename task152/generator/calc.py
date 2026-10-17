"""The golden, computed from the shipped files only: the plant's daily record joined across the two SCADA layouts,
the washes from the work orders, the original modules' wear read on the days after each block's wash at the
inverters' DC input (temperature-corrected), the last twelve months' performance restored by that wear alone and
carried onto the new modules with the side-by-side gain, then the 25 weather years of the long-term record and the
policy's P90. Every rival route the design note names is computed here too, on the same reads."""
from __future__ import annotations

import csv
import datetime as dt
import math
import os

import numpy as np
from openpyxl import load_workbook

import common as K
import records as R

BASE_FROM, BASE_TO = dt.date(2025, 10, 1), dt.date(2026, 9, 30)


class Plant:
    pass


def _d_us(s):
    m, d, y = (int(x) for x in s.split("/"))
    return dt.date(y, m, d)


def read_plant(folder):
    """Daily arrays over every day either SCADA layout holds, inverters in the original order (INV01 to INV16, which
    the crosswalk maps to the block tags; inverter 07's replacement continues its column)."""
    tags = {}
    with open(os.path.join(folder, R.F_TAGS), encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            if row["scada_tag"].startswith("B"):
                orig = row["original_tag"] or "INV07"
                tags[row["scada_tag"]] = int(orig[3:]) - 1
    days = {}

    def slot(d):
        if d not in days:
            days[d] = dict(dc=np.full(K.N_INV, np.nan), ac=np.full(K.N_INV, np.nan), op=np.full(K.N_INV, np.nan),
                           poa=np.nan, tmod=np.nan, tamb=np.nan, exp=np.nan, imp=np.nan)
        return days[d]

    with open(os.path.join(folder, R.F_OLD), encoding="utf-8") as fh:
        rd = csv.reader(fh)
        head = next(rd)
        ix = {c: k for k, c in enumerate(head)}
        for row in rd:
            s = slot(_d_us(row[0]))
            s["poa"] = float(row[ix["POA_Irr_kWh/m2"]])
            s["tmod"] = float(row[ix["TMod_IW_C"]])
            s["tamb"] = float(row[ix["TAmb_C"]])
            for j in range(K.N_INV):
                t = K.OLD_TAGS[j]
                s["dc"][j] = float(row[ix[t + "_DC_kWh"]])
                s["ac"][j] = float(row[ix[t + "_AC_kWh"]])
                s["op"][j] = float(row[ix[t + "_OpHrs"]])
            s["exp"] = float(row[ix["Meter_Export_kWh"]])
            s["imp"] = float(row[ix["Meter_Import_kWh"]])
    with open(os.path.join(folder, R.F_NEW), encoding="utf-8") as fh:
        rd = csv.reader(fh)
        next(rd)
        for stamp, dev, pt, val, unit in rd:
            t = dt.datetime.strptime(stamp, "%Y-%m-%dT%H:%M:%SZ") - R.UTC_OFF
            s = slot(t.date())
            v = float(val)
            if dev == "MET-01":
                if pt == "H_POA":
                    s["poa"] = v
                elif pt == "T_MOD_IW":
                    s["tmod"] = v
                elif pt == "T_AMB":
                    s["tamb"] = v
            elif dev == "PLANT-RM":
                s["exp" if pt == "E_EXP" else "imp"] = v * 1000.0
            else:
                j = tags[dev]
                if pt == "E_DC":
                    s["dc"][j] = v * 1000.0
                elif pt == "E_AC":
                    s["ac"][j] = v * 1000.0
                elif pt == "T_OP":
                    s["op"][j] = v
    P = Plant()
    P.dates = sorted(days)
    P.ix = {d: k for k, d in enumerate(P.dates)}
    P.dc = np.array([days[d]["dc"] for d in P.dates])
    P.ac = np.array([days[d]["ac"] for d in P.dates])
    P.op = np.array([days[d]["op"] for d in P.dates])
    P.poa = np.array([days[d]["poa"] for d in P.dates])
    P.tmod = np.array([days[d]["tmod"] for d in P.dates])
    P.tamb = np.array([days[d]["tamb"] for d in P.dates])
    P.exp = np.array([days[d]["exp"] for d in P.dates])
    P.imp = np.array([days[d]["imp"] for d in P.dates])
    P.year = np.array([d.year for d in P.dates])
    P.month = np.array([d.month for d in P.dates])
    P.dts = np.array(P.dates)
    return P


def read_washes(folder):
    """Each block's full washes from the work orders: (block index, the day the wash was done)."""
    wb = load_workbook(os.path.join(folder, R.F_WO), read_only=True)
    ws = wb.active
    out = []
    first = True
    for row in ws.iter_rows(values_only=True):
        if first:
            first = False
            continue
        no, typ, asset, desc, made, done = row[:6]
        if typ == "CLN" and asset.startswith("BLK"):
            out.append((int(asset[3:]) - 1, done.date()))
    return sorted(out)


def read_lt(folder):
    """The long-term record's annual plane-of-array irradiation, kWh/m2, by local calendar year."""
    tot = {}
    with open(os.path.join(folder, R.F_LT), encoding="utf-8") as fh:
        for line in fh:
            if line.startswith("#") or line.startswith("period_end_utc"):
                continue
            stamp, ghi, dhi, poa = line.rstrip("\n").split(",")
            t = dt.datetime.strptime(stamp, "%Y-%m-%d %H:%M") - R.UTC_OFF - dt.timedelta(hours=1)
            tot[t.year] = tot.get(t.year, 0.0) + float(poa)
    return {y: v / 1000.0 for y, v in tot.items() if 2001 <= y <= 2025}


def read_report(folder):
    """The side-by-side report's twelve-month row: irradiation, table A and table B specific yields, B over A; and
    its monthly rows."""
    import pdfplumber
    import docs as DOC
    with pdfplumber.open(os.path.join(folder, DOC.F_RIG)) as pdf:
        text = "\n".join(pg.extract_text() or "" for pg in pdf.pages)
    import re
    m = re.search(r"Twelve months ([\d.]+) ([\d.]+) ([\d.]+) ([\d.]+)", text)
    months = [(mm.group(1), int(mm.group(2)), float(mm.group(3)), float(mm.group(4)), float(mm.group(5)), float(mm.group(6)))
              for mm in re.finditer(r"([A-Z][a-z]{2}) (20\d\d) ([\d.]+) ([\d.]+) ([\d.]+) ([\d.]+)", text)]
    return dict(poa=float(m.group(1)), a=float(m.group(2)), b=float(m.group(3)), ratio=float(m.group(4)), months=months)


def report_gain(folder):
    """The measured gain: the twelve months' DC energy per kWp of table B over table A, from the report's totals."""
    rep = read_report(folder)
    return rep["b"] / rep["a"]


class Gold:
    pass


def analyse(folder, gain, kwp_old_inv=K.OLD_KWP_INV, gamma=K.GAMMA_OLD, dc_old=K.DC_OLD_KWP, dc_new=K.DC_NEW_KWP):
    P = read_plant(folder)
    washes = read_washes(folder)
    lt = read_lt(folder)
    G = Gold()
    G.P, G.washes, G.lt, G.gain = P, washes, lt, gain
    H = np.array([lt[y] for y in range(2001, 2026)])
    G.H = H
    base = (P.dts >= BASE_FROM) & (P.dts <= BASE_TO)
    y16 = P.year == 2016
    G.base_mask, G.y16 = base, y16
    G.pr = lambda m: float(np.nansum(P.exp[m] - P.imp[m]) / (np.nansum(P.poa[m]) * dc_old))
    G.PR_B, G.PR_16 = G.pr(base), G.pr(y16)
    tc = 1.0 + gamma * (P.tmod - 25.0)
    G.yld = P.dc / (P.poa[:, None] * kwp_old_inv * tc[:, None])
    G.yld_unc = P.dc / (P.poa[:, None] * kwp_old_inv)
    G.yac = P.ac / (P.poa[:, None] * kwp_old_inv * tc[:, None])
    G.full = P.op >= np.nanmax(P.op, axis=1, keepdims=True) - 0.1

    def clean(kmin=1, kmax=3):
        m = np.zeros_like(G.full)
        for b, d in washes:
            for k in range(kmin, kmax + 1):
                i = P.ix.get(d + dt.timedelta(days=k))
                if i is not None:
                    m[i, 4 * b:4 * b + 4] = True
        return m & G.full
    G.clean = clean
    G.CM = clean()

    def level(y, mask, rows):
        mm = mask & rows[:, None]
        return float(np.nanmean(y[mm]))
    G.level = level
    G.L_B, G.L_16 = level(G.yld, G.CM, base), level(G.yld, G.CM, y16)
    G.n_B, G.n_16 = int((G.CM & base[:, None]).sum()), int((G.CM & y16[:, None]).sum())
    G.wear = G.L_B / G.L_16
    G.PR_new = G.PR_B * gain / G.wear
    G.E = G.PR_new * dc_new * H / 1000.0
    G.P50 = float(G.E.mean())
    G.SD = float(G.E.std(ddof=1))
    G.P90 = G.P50 - K.POLICY_Z * G.SD
    G.volume = int(math.floor(G.P90 / K.VOLUME_STEP)) * K.VOLUME_STEP
    Hp90 = float(H.mean() - K.POLICY_Z * H.std(ddof=1))
    G.Hp90 = Hp90
    G.p90_of = lambda prnew: prnew * dc_new * Hp90 / 1000.0
    return G


def routes(G, dc_old=K.DC_OLD_KWP, dc_new=K.DC_NEW_KWP):
    """Every construction the design note names, each a P90 on the same 25 weather years."""
    P, gain = G.P, G.gain
    E = G.p90_of
    r = {}
    r["golden"] = E(G.PR_new)
    base, y16 = G.base_mask, G.y16
    restore = lambda ratio: E(G.PR_B * gain / ratio)

    def lv(y, cm, rows, est=np.nanmean):
        return float(est(y[cm & rows[:, None]]))
    for (a, b) in [(1, 1), (1, 2), (1, 5), (2, 3), (2, 4)]:
        cm = G.clean(a, b)
        r["golden variant: days %d to %d after the wash" % (a, b)] = restore(lv(G.yld, cm, base) / lv(G.yld, cm, y16))
    cm0 = G.clean(0, 3)
    r["golden variant: the wash day included"] = restore(lv(G.yld, cm0, base) / lv(G.yld, cm0, y16))
    CM = G.CM
    rat = []
    for j in range(K.N_INV):
        mb, m16 = CM[:, j] & base, CM[:, j] & y16
        rat.append(np.nanmean(G.yld[mb, j]) / np.nanmean(G.yld[m16, j]))
    r["golden variant: inverter by inverter"] = restore(float(np.mean(rat)))
    r["golden variant: the median"] = restore(lv(G.yld, CM, base, np.nanmedian) / lv(G.yld, CM, y16, np.nanmedian))
    tcm = 1.0 + K.GAMMA_OLD * (P.tmod[:, None] - 25.0)
    num, den = P.dc / tcm, P.poa[:, None] * K.OLD_KWP_INV * np.ones((1, K.N_INV))
    ros = lambda m: float(np.nansum(num[m]) / np.nansum(den[m]))
    r["golden variant: ratio of sums"] = restore(ros(CM & base[:, None]) / ros(CM & y16[:, None]))
    full05 = P.op >= np.nanmax(P.op, axis=1, keepdims=True) - 0.05
    cm5 = G.clean() & full05
    r["golden variant: full operating hours within 0.05 h"] = restore(lv(G.yld, cm5, base) / lv(G.yld, cm5, y16))
    for lab, ref in [("as-new from commercial operation to 2016", P.year <= 2016),
                     ("as-new over the first twelve months", P.dts < K.COD + dt.timedelta(days=366)),
                     ("2016 and 2017 as the reference", (P.year >= 2016) & (P.year <= 2017)),
                     ("commercial operation to 2017 as the reference", P.year <= 2017)]:
        r["golden variant: %s" % lab] = restore(G.L_B / lv(G.yld, CM, ref))
    first_wash = sorted(set(d for b, d in G.washes if d.year == 2015))
    if first_wash:
        fw = (P.dts <= first_wash[-1] + dt.timedelta(days=3))
        r["golden variant: as-new from the handover wash"] = restore(G.L_B / lv(G.yld, CM, fw))
    # each wash's clean level, a line through the one-wash years evaluated at the baseline's middle
    pts = []
    by_wash = {}
    for b, d in G.washes:
        key = (d.year, d.month if d.year > 2015 else 11)
        by_wash.setdefault(key, []).append((b, d))
    for key in sorted(by_wash):
        vals, tt = [], []
        for b, d in by_wash[key]:
            for k in (1, 2, 3):
                i = P.ix.get(d + dt.timedelta(days=k))
                if i is None:
                    continue
                for j in range(4 * b, 4 * b + 4):
                    if G.full[i, j]:
                        vals.append(G.yld[i, j])
                        tt.append((P.dates[i] - K.COD).days / 365.25)
        if vals:
            pts.append((float(np.mean(tt)), float(np.mean(vals)), key))
    G.wash_levels = pts
    G.wash_groups = by_wash
    late = np.array([(t, v) for t, v, k in pts if k[0] >= 2021])
    c = np.polyfit(late[:, 0], late[:, 1], 1)
    tb = (dt.date(2026, 4, 1) - K.COD).days / 365.25
    r["golden variant: one-wash trend at the baseline's middle"] = restore(float(np.polyval(c, tb)) / G.L_16)
    for q in (99, 95):
        fb, f16 = G.full & base[:, None], G.full & y16[:, None]
        r["golden variant: top %d per cent envelope" % (100 - q)] = restore(np.nanpercentile(G.yld[fb], q) / np.nanpercentile(G.yld[f16], q))
    two = np.zeros_like(G.full)
    for b, d in G.washes:
        if d >= dt.date(2024, 12, 1):
            for k in (1, 2, 3):
                i = P.ix.get(d + dt.timedelta(days=k))
                if i is not None:
                    two[i, 4 * b:4 * b + 4] = True
    two &= G.full
    r["golden variant: the last two December washes"] = restore(float(np.nanmean(G.yld[two])) / G.L_16)
    # the datasheet's coefficient applied in its first-order translation form, to 25 degC or to either period's mean
    TTm = np.broadcast_to(P.tmod[:, None], P.dc.shape)
    y1 = G.yld_unc * (1.0 - K.GAMMA_OLD * (TTm - 25.0))
    r["golden variant: first-order translation to 25 degC"] = restore(lv(y1, CM, base) / lv(y1, CM, y16))
    for lab, rows_ in [("the base period's", base), ("2016's", y16)]:
        Tm = float(np.nanmean(TTm[CM & rows_[:, None]]))
        yt = G.yld_unc * (1.0 + K.GAMMA_OLD * (Tm - TTm))
        r["golden variant: first-order translation to %s mean temperature" % lab] = restore(lv(yt, CM, base) / lv(yt, CM, y16))
    # the temperature correction done another way: NOCT from the air temperature, or fitted on the record itself
    Gm = P.poa[:, None] * 1000.0 / np.maximum(P.op, 1.0)
    Tn = P.tamb[:, None] + (45.0 - 20.0) / 800.0 * Gm
    G.yld_noct = G.yld_unc / (1.0 + K.GAMMA_OLD * (Tn - 25.0))
    r["golden variant: NOCT cell temperature from the air temperature"] = restore(lv(G.yld_noct, CM, base) / lv(G.yld_noct, CM, y16))
    r["golden variant: NOCT cell temperature, the median"] = restore(lv(G.yld_noct, CM, base, np.nanmedian) / lv(G.yld_noct, CM, y16, np.nanmedian))
    TT = np.broadcast_to(P.tmod[:, None], P.dc.shape)
    mB, m16 = CM & base[:, None], CM & y16[:, None]
    TB, yB, T6, y6 = TT[mB], G.yld_unc[mB], TT[m16], G.yld_unc[m16]
    kB, k6 = np.polyfit(TB, yB, 1), np.polyfit(T6, y6, 1)
    pv = np.polyval
    r["golden variant: regression on module temperature in each period, at 25 degC"] = restore(pv(kB, 25.0) / pv(k6, 25.0))
    r["golden variant: regression in each period, at the base period's mean temperature"] = restore(pv(kB, TB.mean()) / pv(k6, TB.mean()))
    X = np.column_stack([np.r_[np.ones(len(yB)), np.zeros(len(y6))], np.r_[np.zeros(len(yB)), np.ones(len(y6))], np.r_[TB, T6]])
    cl = np.linalg.lstsq(X, np.log(np.r_[yB, y6]), rcond=None)[0]
    r["golden variant: one temperature slope for both periods, log yield"] = restore(math.exp(cl[0] - cl[1]))
    cx = np.linalg.lstsq(X, np.r_[yB, y6], rcond=None)[0]
    r["golden variant: one temperature slope for both periods, at 25 degC"] = restore((cx[0] + 25.0 * cx[2]) / (cx[1] + 25.0 * cx[2]))
    gh = k6[0] / pv(k6, 25.0)
    r["golden variant: the coefficient fitted on 2016"] = restore(float(np.mean(yB / (1 + gh * (TB - 25.0))) / np.mean(y6 / (1 + gh * (T6 - 25.0)))))
    TnB, Tn6 = Tn[mB], Tn[m16]
    r["golden variant: regression on NOCT cell temperature in each period, at 25 degC"] = restore(
        pv(np.polyfit(TnB, yB, 1), 25.0) / pv(np.polyfit(Tn6, y6, 1), 25.0))
    G.temp_fit = dict(slope_base=float(kB[0] / pv(kB, 25.0)), slope_2016=float(gh), t_base=float(TB.mean()), t_2016=float(T6.mean()),
                      tn_base=float(TnB.mean()), tn_2016=float(Tn6.mean()))
    # the season matched as well as corrected: December 2025 against 2016's winter washes
    wm = np.zeros_like(G.full, dtype=int)
    for b, d in G.washes:
        for k in (1, 2, 3):
            i = P.ix.get(d + dt.timedelta(days=k))
            if i is not None:
                wm[i, 4 * b:4 * b + 4] = d.month
    for lab, ms in [("2016's January and November washes", (1, 11)), ("January 2016", (1,)), ("November 2016", (11,))]:
        ref = CM & np.isin(wm, ms) & y16[:, None]
        r["golden variant: December 2025 against %s" % lab] = restore(G.L_B / float(np.nanmean(G.yld[ref])))
        if len(ms) == 2:
            G.winter_ref = ref
    # the baseline window
    lastw = {}
    for t, v, k in pts:
        if k[0] >= 2020:
            lastw[k[0]] = v
    for lab, ys in [("2025", [2025]), ("2024 and 2025", [2024, 2025]), ("2021 to 2025", [2021, 2022, 2023, 2024, 2025]),
                    ("2022 to 2025", [2022, 2023, 2024, 2025]), ("2023 to 2025", [2023, 2024, 2025])]:
        lev = [G.pr(P.year == y) / (lastw[y] / G.L_16) for y in ys]
        r["golden variant: kept losses of %s" % lab] = E(float(np.mean(lev)) * gain)
    m24 = (P.dts >= dt.date(2024, 10, 1)) & (P.dts <= BASE_TO)
    r["golden variant: the last 24 months"] = E(G.pr(m24) * gain / G.wear)
    # rivals
    r["R0 the last twelve months carried with the gain"] = E(G.PR_B * gain)
    r["R0b the last twelve months on the new nameplate"] = E(G.PR_B)
    r["R1 the first year carried with the gain"] = E(G.PR_16 * gain)
    r["R1b the first year on the new nameplate"] = E(G.PR_16)
    fB, f16 = G.full & base[:, None], G.full & y16[:, None]
    den = P.poa[:, None] * K.OLD_KWP_INV * (1.0 + K.GAMMA_OLD * (P.tmod[:, None] - 25.0)) * np.ones((1, K.N_INV))
    den_u = P.poa[:, None] * K.OLD_KWP_INV * np.ones((1, K.N_INV))
    ratio = lambda num, dd, m1, m0: (np.nansum(num[m1]) / np.nansum(dd[m1])) / (np.nansum(num[m0]) / np.nansum(dd[m0]))
    allB = np.ones_like(G.full) & base[:, None]
    all16 = np.ones_like(G.full) & y16[:, None]
    r["R2 the DC-side decline on available days"] = restore(ratio(P.dc, den, fB, f16))
    r["R2 variant: the DC-side decline on every day"] = restore(ratio(P.dc, den, allB, all16))
    r["R2 variant: the DC-side decline, no temperature correction"] = restore(ratio(P.dc, den_u, fB, f16))
    r["R2b the decline after washes at the inverters' AC output"] = restore(lv(G.yac, CM, base) / lv(G.yac, CM, y16))
    yac_u = P.ac / (P.poa[:, None] * K.OLD_KWP_INV)
    r["R2b variant: the decline after washes at the AC output, no temperature correction"] = restore(lv(yac_u, CM, base) / lv(yac_u, CM, y16))
    r["R2b variant: the decline after washes at the AC output, no temperature correction, the median"] = restore(
        lv(yac_u, CM, base, np.nanmedian) / lv(yac_u, CM, y16, np.nanmedian))
    r["R2b variant: the decline after washes, no temperature correction"] = restore(lv(G.yld_unc, CM, base) / lv(G.yld_unc, CM, y16))
    r["R2b variant: the decline after washes, no temperature correction, the median"] = restore(
        lv(G.yld_unc, CM, base, np.nanmedian) / lv(G.yld_unc, CM, y16, np.nanmedian))
    rosu = lambda m: float(np.nansum(P.dc[m]) / np.nansum(den_u[m]))
    r["R2b variant: the decline after washes, no temperature correction, ratio of sums"] = restore(rosu(CM & base[:, None]) / rosu(CM & y16[:, None]))
    r["R2b variant: December 2025 against 2016's winter washes, no temperature correction"] = restore(
        lv(G.yld_unc, CM, base) / float(np.nanmean(G.yld_unc[G.winter_ref])))
    # the plant as a whole in the week after each wash's last block, read at the revenue meter
    wk_b, wk_16 = np.zeros(len(P.dates), bool), np.zeros(len(P.dates), bool)
    for key, items in by_wash.items():
        last = max(d for b, d in items)
        for k in range(1, 8):
            i = P.ix.get(last + dt.timedelta(days=k))
            if i is not None and base[i]:
                wk_b[i] = True
            elif i is not None and y16[i]:
                wk_16[i] = True
    prn = (P.exp - P.imp) / (P.poa * dc_old)
    r["R2e the plant's net ratio in the week after each wash, no temperature correction"] = restore(
        float(np.nanmean(prn[wk_b])) / float(np.nanmean(prn[wk_16])))
    G.plant_week = (wk_b, wk_16)
    yrs = list(range(2016, 2026))
    yy = np.array([np.nanmean(G.yld[P.year == y][G.full[P.year == y]]) for y in yrs])
    G.yearly_yield = yy
    yoy = yy[1:] / yy[:-1] - 1.0
    age = (dt.date(2026, 4, 1) - dt.date(2016, 7, 1)).days / 365.25
    r["R2c the year-on-year median rate times the plant's age"] = restore(1.0 + float(np.median(yoy)) * age)
    s5 = np.polyfit(np.arange(2021, 2026), yy[-5:] / yy[0], 1)[0]
    r["R2c the last five years' trend times the plant's age"] = restore(1.0 + s5 * age)
    s10 = np.polyfit(np.arange(2016, 2026), yy / yy[0], 1)[0]
    r["R2c the whole record's trend times the plant's age"] = restore(1.0 + s10 * age)
    agec = (dt.date(2026, 4, 1) - K.COD).days / 365.25
    r["R-warranty the original modules' warranted decline"] = restore(1.0 - 0.025 - 0.007 * (agec - 1.0))
    ages = [("at the base period's middle", agec), ("at the sale", (dt.date(2026, 10, 6) - K.COD).days / 365.25)]
    for rate in (0.005, 0.008, 0.010, 0.015):
        for lab, a_ in ages:
            r["R-rate %.1f per cent a year, straight line, %s" % (100 * rate, lab)] = restore(1.0 - rate * a_)
            r["R-rate %.1f per cent a year on the warranty's schedule, %s" % (100 * rate, lab)] = restore(1.0 - 0.025 - rate * (a_ - 1.0))
    windows = {}
    for a in range(2016, 2026):
        for b in range(a, 2026):
            m = (P.year >= a) & (P.year <= b)
            windows[(a, b)] = (E(G.pr(m) * gain), E(G.pr(m)))
    G.windows = windows
    for lab, a0, b0 in [("the whole record", K.COD, BASE_TO), ("2016 to 2025", dt.date(2016, 1, 1), dt.date(2025, 12, 31)),
                        ("the original SCADA's years", K.COD, K.SCADA_CUTOVER - dt.timedelta(days=1)),
                        ("the replacement SCADA's years", K.SCADA_CUTOVER, BASE_TO),
                        ("the one-wash years", dt.date(2021, 1, 1), BASE_TO), ("the last five years", dt.date(2021, 1, 1), dt.date(2025, 12, 31)),
                        ("the last three years", dt.date(2023, 1, 1), dt.date(2025, 12, 31)), ("2025", dt.date(2025, 1, 1), dt.date(2025, 12, 31)),
                        ("the first twelve months", K.COD, K.COD + dt.timedelta(days=364)), ("the first three years", dt.date(2016, 1, 1), dt.date(2018, 12, 31)),
                        ("the six-wash years", dt.date(2016, 1, 1), dt.date(2018, 12, 31)), ("the years before one wash a year", dt.date(2016, 1, 1), dt.date(2020, 12, 31))]:
        m = (P.dts >= a0) & (P.dts <= b0)
        r["W %s with the gain" % lab] = E(G.pr(m) * gain)
        r["W %s on the new nameplate" % lab] = E(G.pr(m))
    return r
