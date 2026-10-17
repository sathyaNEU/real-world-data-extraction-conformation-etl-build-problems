"""task152 checks: the three rules executed on the clean truth, the ladder and every rival route, the corridor of
the golden construction, every semantic fork row, the closure axes that can move, both asks with their stops and
devices, the separation line, the realism sweep, the input gates and the distractors. Run by build.py; exits on the
first report with any failure."""
from __future__ import annotations

import csv
import datetime as dt
import json
import math
import os
import re
import shutil
import tempfile

import numpy as np
import pdfplumber

import asks as AS
import calc as CA
import common as K
import deliver as DL
import docs as DOC
import records as R
import world as WD

RESULTS = []


def ok(label, cond, detail=""):
    RESULTS.append((label, bool(cond), detail))


def floor_k(x):
    return int(math.floor(x / 1000.0)) * 1000


def run(target, golden, B=None, G=None, meta_path=None):
    W = B.W
    A = G["A"]
    r = G["routes"]
    gold = r["golden"]
    vol = G["volume"]
    gfloor = floor_k(gold)
    p50f = lambda p90: p90 * A.P50 / A.P90         # the P50 that goes with a P90 on the same weather years
    nearest = lambda x, s: int(math.floor(x / s + 0.5)) * s

    # ------------------------------------------------------------------ the call and its bins
    ok("CALL volume 70,000 MWh", vol == 70000 and gfloor == 70000, "P90 %.1f" % gold)
    ok("CALL P90 mid-bin (floor clearance >= 0.3)", K.floor_clearance(gold, 1000) >= 0.3,
       "clearance %.2f" % K.floor_clearance(gold, 1000))
    ok("CALL P50 to the nearest 1,000 is 73,000, at least 400 MWh from both rounding edges",
       G["P50_r"] == 73000 and min(A.P50 - 72500, 73500 - A.P50) >= 400, "P50 %.1f" % A.P50)
    ok("CALL standard deviation to the nearest 100 is 1,900, at least 20 MWh from its edges (the population SD too)",
       G["SD_r"] == 1900 and min(abs(A.SD - 1850), abs(A.SD - 1950)) >= 20
       and min(abs(A.SD * math.sqrt(24 / 25) - 1850), abs(A.SD * math.sqrt(24 / 25) - 1950)) >= 20,
       "SD %.1f, population %.1f" % (A.SD, A.SD * math.sqrt(24 / 25)))
    ok("CALL P90 = P50 - 1.2816 SD over the 25 weather years", abs(A.P90 - (A.P50 - 1.2816 * A.SD)) < 1e-6)
    ok("CALL 25 weather years, 2001 to 2025", len(A.H) == 25 and sorted(A.lt) == list(range(2001, 2026)))

    # ------------------------------------------------------------------ the three rules
    nB, n16 = A.n_B, A.n_16
    ok("R1a the decisive computation reads the inverter-days after each block's wash in both SCADA layouts",
       nB >= 40 and n16 >= 250 and A.P.dates[0] == K.COD, "%d base-period and %d 2016 inverter-days" % (nB, n16))
    bad = []
    for f in sorted(os.listdir(target)):
        if f.endswith(".csv"):
            with open(os.path.join(target, f), encoding="utf-8") as fh:
                head = next(ln for ln in fh if not ln.startswith("#"))
            if re.search(r"soil|wear|degrad|clean|perf|ratio|pr_|loss", head, re.I):
                bad.append(f)
    ok("R1b no shipped column holds a soiling, wear, loss or performance-ratio figure", not bad, str(bad))
    # R2: each Layer 1 item mishandled leaves the call where it is
    P = A.P
    stamps = {ln.split(",", 1)[0] for ln in open(os.path.join(target, R.F_NEW), encoding="utf-8") if ln[0] == "2"}
    ok("R2a the replacement SCADA's day_start_utc read as a UTC date gives the same days (every stamp is 07:00Z)",
       all(x.endswith("T07:00:00Z") for x in stamps), "%d day stamps" % len(stamps))
    tmp = tempfile.mkdtemp(dir=os.path.dirname(os.path.abspath(target)), suffix=".tmp")
    try:
        # INV07's replacement left unmapped: its days drop out
        for f in os.listdir(target):
            shutil.copy(os.path.join(target, f), os.path.join(tmp, f))
        rows = list(csv.reader(open(os.path.join(target, R.F_NEW), encoding="utf-8")))
        K.write_csv(os.path.join(tmp, R.F_NEW), rows[0], [x for x in rows[1:] if x[1] != K.NEW_TAG_07R])
        A2 = CA.analyse(tmp, G["gain"])
        ok("R2b inverter 07's replacement unmapped (B2-INV3R dropped): same volume", A2.volume == vol, "%.0f" % A2.P90)
    finally:
        shutil.rmtree(tmp)
    # the wash day is the work orders' Completed date: each block's yield steps up across it
    steps25, steps16 = [], []
    for b, d in A.washes:
        if d.year == 2016 or (CA.BASE_FROM <= d <= CA.BASE_TO):
            bef = [A.P.ix[d - dt.timedelta(days=k)] for k in (1, 2, 3) if d - dt.timedelta(days=k) in A.P.ix]
            aft = [A.P.ix[d + dt.timedelta(days=k)] for k in (1, 2, 3) if d + dt.timedelta(days=k) in A.P.ix]
            yb_ = np.nanmean(np.where(A.full[bef, 4 * b:4 * b + 4], A.yld[bef, 4 * b:4 * b + 4], np.nan))
            ya_ = np.nanmean(np.where(A.full[aft, 4 * b:4 * b + 4], A.yld[aft, 4 * b:4 * b + 4], np.nan))
            (steps16 if d.year == 2016 else steps25).append(ya_ / yb_ - 1.0)
    ok("R2c the wash day is the work orders' Completed date: every block's DC yield steps up across it in December "
       "2025 (by 2 per cent or more) and across 2016's washes on average", len(steps25) == 4 and min(steps25) > 0.02
       and float(np.mean(steps16)) > 0.0, "December 2025 %s; 2016 mean %+.4f" % (" ".join("%+.3f" % x for x in steps25), float(np.mean(steps16))))
    # the long-term record grouped by UTC year
    tot = {}
    with open(os.path.join(target, R.F_LT), encoding="utf-8") as fh:
        for line in fh:
            if line[0] in "#p":
                continue
            st, g_, d_, poa = line.rstrip("\n").split(",")
            tot[int(st[:4])] = tot.get(int(st[:4]), 0.0) + float(poa)
    Hu = np.array([tot[y] / 1000.0 for y in range(2001, 2026)])
    pu = A.PR_new * K.DC_NEW_KWP * (Hu.mean() - 1.2816 * Hu.std(ddof=1)) / 1000.0
    ok("R2d the long-term record grouped by UTC year: same volume", floor_k(pu) == vol, "%.0f" % pu)
    rep = CA.read_report(target)
    gm = float(np.mean([m[5] for m in rep["months"]]))
    pm = A.p90_of(A.PR_B * gm / A.wear)
    ok("R2e the report's gain read as the mean of its monthly ratios: same volume", floor_k(pm) == vol, "%.0f" % pm)
    lost = W.old_lost
    ok("R2f the original SCADA's lost days (August 2018) touch neither 2016 nor the base period",
       all(d.year == 2018 for d in lost) and not any(d in P.ix for d in lost))
    ok("R2g the base period lies wholly after the station service moved behind the revenue meter",
       CA.BASE_FROM > K.STATION_SERVICE_END)
    # R3a: the analyst test on the clean truth
    yrs = np.array([d.year for d in W.plant_days])
    dts = np.array(W.plant_days)
    bmask = (dts >= CA.BASE_FROM) & (dts <= CA.BASE_TO)
    pr_true = float((W.del_d[bmask] - W.rec_d[bmask]).sum() / (W.poa_d[bmask].sum() * K.DC_OLD_KWP))
    Ht = np.array([W.wx[y]["poa"].sum() / 1000.0 for y in range(2001, 2026)])
    hp90 = Ht.mean() - 1.2816 * Ht.std(ddof=1)
    r0_true = pr_true * W.rig_gain * K.DC_NEW_KWP * hp90 / 1000.0
    tcw = 1.0 + K.GAMMA_OLD * (W.tmod_w - 25.0)
    yt = W.dc_d / (W.poa_d[:, None] * K.OLD_KWP_INV * tcw[:, None])
    full = W.op_hours >= W.op_hours.max(axis=1, keepdims=True) - 0.1
    cmt = np.zeros_like(full)
    for no, b, d in W.wash_events:
        for k in (1, 2, 3):
            i = W.day_index.get(d + dt.timedelta(days=k))
            if i is not None:
                cmt[i, 4 * b:4 * b + 4] = True
    cmt &= full
    wear_t = yt[cmt & bmask[:, None]].mean() / yt[cmt & (yrs == 2016)[:, None]].mean()
    gold_true = pr_true * W.rig_gain / wear_t * K.DC_NEW_KWP * hp90 / 1000.0
    ok("R3a analyst test on the clean truth: the last twelve months carried with the gain files %d, not %d" % (
        floor_k(r0_true), vol), floor_k(r0_true) != vol and abs(r0_true - r["R0 the last twelve months carried with the gain"]) < 1000,
       "%.0f" % r0_true)
    ok("R3a the golden construction on the clean truth files the same volume", floor_k(gold_true) == vol, "%.0f" % gold_true)
    ok("R3b the default path with only the decisive step changed files the golden",
       floor_k(A.p90_of(A.PR_B * G["gain"] / A.wear)) == vol)
    for k in ("R1 the first year carried with the gain", "R2 the DC-side decline on available days",
              "R2b the decline after washes at the inverters' AC output"):
        ok("R3b a lower rung's single correction misses: %s" % k, floor_k(r[k]) != vol, "%.0f" % r[k])
    # R3c and the HOW sentences: what the filed sentences give when executed literally
    lit = A.p90_of(A.PR_B * G["gain"])
    ok("R3c the filed sentences executed literally (the base period, the report's gain, the scope's capacity, the "
       "policy's P90) give no golden: their only literal product is %d" % floor_k(lit), floor_k(lit) != vol)
    texts = {}
    for f in sorted(os.listdir(target)):
        p = os.path.join(target, f)
        if f.endswith(".pdf"):
            with pdfplumber.open(p) as pdf:
                texts[f] = "\n".join(pg.extract_text() or "" for pg in pdf.pages)
        elif f.endswith(".md"):
            texts[f] = open(p, encoding="utf-8").read()
        elif f.endswith(".docx"):
            import docx
            texts[f] = "\n".join(pp.text for pp in docx.Document(p).paragraphs)
    alltext = "\n".join(texts.values())
    ok("R3c no document states a 2027 volume, P50, wear or performance figure for the plant",
       not re.search(r"(2027[^.]{0,60}(MWh|P50|P90))|(\bwear\b)|(performance ratio)|(\bdegradation\b)", alltext, re.I),
       re.findall(r"[^.]{0,40}(?:2027[^.]{0,60}(?:MWh|P50|P90)|\bwear\b|performance ratio|\bdegradation\b)", alltext, re.I)[:3])
    ok("R3c no plant: no 7-day run of any inverter's daily DC energy repeats elsewhere", _no_repeats(W))

    # ------------------------------------------------------------------ the ladder and every rival
    rivals = {k: v for k, v in r.items() if k.startswith(("R0", "R1", "R2", "R-"))}
    for k, v in rivals.items():
        out_by = max(vol - v, v - (vol + 1000))
        ok("LADDER %s misses the golden's bin by at least 300 MWh" % k, floor_k(v) != vol and out_by >= 300,
           "%.0f (%+.2f bins, %.0f outside)" % (v, (v - gold) / 1000, out_by))
    ok("LADDER rung 0 is at least 5 bins under and rung 1 at least 5 over",
       r["R0 the last twelve months carried with the gain"] < gold - 5000 and r["R1 the first year carried with the gain"] > gold + 5000)
    ok("LADDER the natural expert's rival (the decline after washes at AC) is at least 1.4 bins off",
       abs(r["R2b the decline after washes at the inverters' AC output"] - gold) >= 1400)
    natural = ["W the whole record with the gain", "W 2016 to 2025 with the gain", "W the original SCADA's years with the gain",
               "W the replacement SCADA's years with the gain", "W the one-wash years with the gain",
               "W the last five years with the gain", "W the last three years with the gain", "W 2025 with the gain",
               "W the first twelve months with the gain", "W the first three years with the gain",
               "W the six-wash years with the gain", "W the years before one wash a year with the gain"]
    natural += [k.replace("with the gain", "on the new nameplate") for k in natural]
    for k in natural:
        ok("WINDOW %s misses" % k[2:], floor_k(r[k]) != vol, "%.0f" % r[k])
    landing = sorted("%d-%d" % ab for ab, (wg, wn) in A.windows.items() if floor_k(wg) == vol or floor_k(wn) == vol)
    ok("WINDOW the windows that land are unmarked middle spans", all(x in ("2016-2021", "2017-2020", "2018-2019", "2017-2021",
                                                                        "2018-2020", "2016-2022") for x in landing), landing)
    # the side-by-side reference table against the plant
    base_dc = np.nansum(A.P.dc[A.base_mask] * A.full[A.base_mask]) / (np.nansum(A.full[A.base_mask]) / 16.0) / K.DC_OLD_KWP
    plant_kwp = np.nansum(A.P.dc[A.base_mask][A.full[A.base_mask]]) / (A.full[A.base_mask].sum() * K.OLD_KWP_INV) * 365.0
    rig_route = A.p90_of(A.PR_B * G["gain"] / (plant_kwp / rep["a"]))
    ok("RIVAL the plant's DC energy per kWp over the report's table A, twelve months, restored: misses", floor_k(rig_route) != vol,
       "%.0f" % rig_route)
    best = []
    for (mon, y, poa, a, b, ratio) in rep["months"]:
        m = {"Jan": 1, "Feb": 2, "Mar": 3, "Apr": 4, "May": 5, "Jun": 6, "Jul": 7, "Aug": 8, "Sep": 9, "Oct": 10, "Nov": 11, "Dec": 12}[mon]
        mm = (A.P.year == y) & (A.P.month == m)
        f_ = A.full & mm[:, None]
        pk = np.nansum(A.P.dc[f_]) / (f_.sum() * K.OLD_KWP_INV) * mm.sum()
        best.append(pk / a)
    rb = A.p90_of(A.PR_B * G["gain"] / max(best))
    ok("RIVAL the plant over table A in its best month, restored: misses", floor_k(rb) != vol, "%.0f" % rb)
    tres = float(np.mean([float(p[7]) for p in DOC.TRESCO_PLANTS])) / 100.0
    td = A.PR_B * (1 + tres) * K.DC_OLD_KWP * A.Hp90 / 1000.0
    ok("DISTRACTOR Tresco's mean repowering uplift on the base period misses", floor_k(td) != vol, "%.0f" % td)

    # ------------------------------------------------------------------ the corridor of the golden construction
    corridor = [k for k in r if k.startswith("golden variant")]
    defensible = [k for k in corridor if "wash day included" not in k and "top 5 per cent" not in k and "kept losses" not in k
                  and "last 24 months" not in k]
    for k in defensible:
        v = r[k]
        m90 = min(v - vol, vol + 1000 - v)
        m50 = min(p50f(v) - (G["P50_r"] - 500), G["P50_r"] + 500 - p50f(v))
        ok("C3 %s: same volume, P50 and standard deviation, each at least 200 MWh inside its bin" % k[15:],
           floor_k(v) == vol and nearest(p50f(v), 1000) == G["P50_r"] and nearest(p50f(v) * A.SD / A.P50, 100) == G["SD_r"]
           and m90 >= 200 and m50 >= 200, "P90 %.0f (%.0f inside), P50 %.0f (%.0f inside)" % (v, m90, p50f(v), m50))
    for k in [k for k in corridor if "kept losses" in k or "last 24 months" in k]:
        ok("SF7 base period %s: same volume (the policy's base period rules it out for the P50)" % k[15:],
           floor_k(r[k]) == vol, "%.0f" % r[k])
    # one standard error on the wear, the gain and the base period's ratio
    yb = A.yld[A.CM & A.base_mask[:, None]]
    y6 = A.yld[A.CM & A.y16[:, None]]
    se_w = A.wear * math.sqrt((yb.std(ddof=1) / yb.mean()) ** 2 / len(yb) + (y6.std(ddof=1) / y6.mean()) ** 2 / len(y6))
    ratios = np.array([m[5] for m in rep["months"]])
    se_g = ratios.std(ddof=1) / math.sqrt(len(ratios)) * 0.3          # months move together; a third of the naive
    for sw in (-1, 1):
        for sg in (-1, 1):
            v = A.p90_of(A.PR_B * (G["gain"] + sg * se_g) / (A.wear + sw * se_w))
            ok("SIGMA wear %+d s.e., gain %+d s.e.: same volume" % (sw, sg), floor_k(v) == vol, "%.0f" % v)

    # ------------------------------------------------------------------ the semantic fork table
    pol = " ".join(DOC.sentences([p for _, _, _, ps in DOC.POLICY for p in ps]))
    ok("SF1 the hedge year is the calendar year (policy 2.1)", "The hedge year is the calendar year" in pol)
    e10 = float(np.percentile(A.E, 10))
    ok("SF2 the P90 is the policy's P50 less 1.2816 standard deviations (4.2); the empirical tenth percentile files %d"
       % floor_k(e10), "P50 less 1.2816 standard deviations" in pol, "%.0f" % e10)
    pop = A.P50 - 1.2816 * float(A.E.std(ddof=0))
    ok("SF3 the population standard deviation: same volume and same rounded SD", floor_k(pop) == vol and
       nearest(float(A.E.std(ddof=0)), 100) == G["SD_r"], "%.0f" % pop)
    ok("SF4 the full calendar years of the record (policy 4.2): 2026 is not in the record", 2026 not in A.lt)
    gross = A.p90_of(float(np.nansum(A.P.ac[A.base_mask]) / (np.nansum(A.P.poa[A.base_mask]) * K.DC_OLD_KWP)) * G["gain"] / A.wear)
    ok("SF5 net energy at the delivery point (policy 3.1): the inverters' gross AC would file another volume",
       floor_k(gross) != vol and "net energy at its delivery point" in pol, "%.0f" % gross)
    sc = " ".join(DOC.scope_sentences())
    ok("SF6 the plant in 2027: the scope's 37,324.8 kWp, the inverters kept, the O&M agreement to 2027",
       "37,324.8 kWp" in sc and "sixteen central inverters" in sc and "December 31, 2027" in sc)
    ok("SF7 the base period is filed (policy 4.4)", "twelve complete months before the sale as the base period" in pol)
    ok("SF8 the as-new reference: 2016, commercial operation to 2016 or the handover wash give the same volume",
       floor_k(r["golden variant: as-new from commercial operation to 2016"]) == vol and
       floor_k(r["golden variant: as-new from the handover wash"]) == vol)
    ok("SF9 the days after the wash: 1 to 2, 1 to 5 and 2 to 4 give the same volume",
       all(floor_k(r["golden variant: days %s after the wash" % x]) == vol for x in ("1 to 2", "1 to 5", "2 to 4")))
    tamb = A.P.dc / (A.P.poa[:, None] * K.OLD_KWP_INV)
    ok("SF10 the temperature basis is the module's own (T_MOD_IW, the datasheet's coefficient for Pmax)",
       "T_MOD_IW" in texts[DOC.F_README] and "-0.42 %/degC" in texts[DOC.F_SHEETS])
    unc = [k for k in r if k.startswith("R2b variant") and "no temperature correction" in k] + \
        ["R2e the plant's net ratio in the week after each wash, no temperature correction"]
    ok("SF10b an uncorrected comparison after the washes files another volume at the DC input, the AC output or the meter "
       "(the base period's post-wash days run %.0f degC cooler than 2016's)" % (A.temp_fit["t_2016"] - A.temp_fit["t_base"]),
       len(unc) >= 6 and all(floor_k(r[k]) != vol for k in unc), " ".join("%.0f" % r[k] for k in unc))
    g135 = A.p90_of(A.PR_B * (1 + G["gain_pct"] / 100.0) / A.wear)
    ok("SF11 the gain as the report's totals, its stated per cent or its monthly mean: same volume",
       floor_k(g135) == vol and floor_k(pm) == vol)
    rig_flat = re.sub(r"\s+", " ", texts[DOC.F_RIG])
    ok("SF12 the report's test conditions: modules measured from first field exposure (the new modules' first year "
       "inside the twelve months), table A's from sealed cartons since 2015",
       "measured from first field exposure" in rig_flat and "sealed cartons since delivery in 2015" in rig_flat)
    ok("SF13 the new nameplate is the scope's (no flash list ships)", not any("flash" in f for f in os.listdir(target)))
    peak, clip_hours, clip_share = 0.0, 0, 0.0
    for y in range(2001, 2026):
        w = W.wx[y]
        ac_new = (K.NEW_KWP_INV * (w["poa"] + K.BIFACIALITY * w["rear"]) / 1000.0 * WD.ETA_REF_NEW *
                  (1 + K.GAMMA_NEW * (w["tmod"] - 27.0)) * W.eta[-1].max())
        peak = max(peak, float(ac_new.max()))
        clip_hours += int((ac_new > K.INV_AC_KW).sum())
        clip_share = max(clip_share, float(np.clip(ac_new - K.INV_AC_KW, 0, None).sum() / ac_new.sum()))
    oldpk = W.ac_h.max()
    ok("SF14 no clipping that moves a figure: the original array's highest hour under 92 per cent of rating; the "
       "repowered array reaches its rating in at most three cloud-edge hours of the 25 years and loses under 0.001 per "
       "cent of any year's output", oldpk < 0.92 * K.INV_AC_KW and clip_hours <= 3 and clip_share < 1e-5,
       "%.3f and %.4f of rating, %d hours, worst year %.6f per cent" % (oldpk / K.INV_AC_KW, peak / K.INV_AC_KW, clip_hours, 100 * clip_share))
    # SF15: the SCADA's revenue meter against the district's 0417, month by month in the base period
    Bk = G["askB"][0]
    dist = {}
    for sp, ch, d, k, rt, v in Bk.mdm:
        if sp == "0417" and CA.BASE_FROM <= d <= dt.date(2025, 12, 31):
            key = (d.year, d.month)
            dist.setdefault(key, {}).setdefault((ch, d), []).append((v, k))
    worst = 0.0
    for key, items in dist.items():
        net_d = sum((max(vs)[1] if ch == "REC" else -max(vs)[1]) for (ch, d), vs in items.items())
        m = (A.P.year == key[0]) & (A.P.month == key[1])
        net_s = float(np.nansum(A.P.exp[m] - A.P.imp[m]))
        worst = max(worst, abs(net_s / net_d - 1.0))
    ok("SF15 the SCADA's revenue meter and the district's service point 0417 agree within 0.1 per cent each month",
       worst < 0.001, "worst %.5f" % worst)
    ok("SF16 the record is the provider's site-adapted series (its header)", "site-adapted" in open(os.path.join(target, R.F_LT)).read(600))
    ok("SF17 the base period lies in the replacement layout and 2016 in the original", CA.BASE_FROM > K.SCADA_CUTOVER)
    ok("SF18 inverter 07's replacement continues its column (crosswalk)", any(K.NEW_TAG_07R in ln for ln in open(os.path.join(target, R.F_TAGS))))
    near = nearest(gold, 1000)
    ok("SF19 rounded to the nearest 1,000 rather than down would file %d; the policy says down" % near,
       near != vol and "rounded down to the nearest" in pol)
    rp = floor_k(G["P50_r"] - 1.2816 * G["SD_r"])
    ok("SF21 the volume rounded down from the unrounded P90; the rounded P50 less 1.2816 rounded standard deviations "
       "gives the same volume", rp == vol, "%d" % rp)
    ok("SF20 the record carries irradiance only, so each weather year's output is the level times its irradiation",
       "poa_wm2" in open(os.path.join(target, R.F_LT)).read(2000) and "tamb" not in open(os.path.join(target, R.F_LT)).read(2000))

    # ------------------------------------------------------------------ closure axes that can move
    nb_days = int(A.base_mask.sum())
    ok("AX12 no missing SCADA day in the base period", nb_days == 365 and not np.isnan(A.P.poa[A.base_mask]).any())
    ok("AX15 the two layouts meet on March 3, 2020 with no day in both",
       sum(1 for _ in csv.reader(open(os.path.join(target, R.F_OLD)))) - 1 + 3 == (K.SCADA_CUTOVER - K.COD).days)
    ok("AX22 nothing in the record changes the plant's kept losses inside 2027: no wash schedule beyond Schedule 2",
       "once in each Contract Year" in " ".join(DOC.OM_S2))
    base_washes = [(b, d) for b, d in A.washes if CA.BASE_FROM <= d <= CA.BASE_TO]
    ok("AX23 the hedge year's dust cycle is the base period's: the base period's only wash is December 2025's, block by "
       "block, Schedule 2 washes in December each Contract Year and the scope keeps the December 2026 wash",
       sorted(b for b, d in base_washes) == [0, 1, 2, 3] and all((d.year, d.month) == (2025, 12) for b, d in base_washes)
       and "block by block in December" in " ".join(DOC.OM_S2) and "December wash goes ahead" in sc)
    whole = 0
    for i in np.where(A.base_mask)[0]:
        dd = A.P.dates[i]
        ages = [min([(dd - d).days for bb, d in A.washes if bb == b and d <= dd] or [999]) for b in range(4)]
        whole += int(min(ages) >= 1 and max(ages) <= 10)
    ok("AX24 no base-period day has every block within ten days of its wash (a block a week), so the plant as a whole "
       "is never read just after a wash in the base period", whole == 0, "%d days" % whole)

    # ------------------------------------------------------------------ Ask A
    Aa, ga, sa = G["askA"]
    for k, st in sa.items():
        miss = sum(1 for t in ga if round(st[t]["avail"], 2) != round(ga[t]["avail"], 2))
        need = {"S1": 16, "S2": 16, "S3": 16, "S4": 14, "S5": 14, "S6": 12, "S7": 1, "S8": 16, "S9": 16, "S10": 4,
                "S11": 4}[k.split()[0]]
        ok("ASK A stop %s misses %d of 16" % (k, miss), miss >= need)
    s4 = re.sub(r"\s+", " ", " ".join(DOC.OM_S4))
    ok("ASK A Schedule 4 states its readings: the formula with its denominator, zero AC energy, a part-hour of "
       "operation, a part-covered Excluded Hour (S10 and S11 are the readings it closes)",
       "A = 100 x (1 - U / (P - X))" in s4 and "is zero" in s4 and "however little of the hour it ran" in s4
       and "only in part" in s4)
    same_a = []
    for kw in (dict(zero_below=1.0), dict(zero_below=5.0), dict(grid_pad=0), dict(grid_pad=30)):
        st_ = AS.availability(Aa, **kw)
        same_a.append(all(round(st_[t]["avail"], 2) == round(ga[t]["avail"], 2) for t in ga))
    small = min(Aa.hours[t][h][0] for t in Aa.hours for h in Aa.prod if h in Aa.hours[t] and Aa.hours[t][h][0] > 0)
    ok("ASK A the readings that change nothing: an hour under 1 or 5 kWh read as none, the district's restoration "
       "read with 0 or 30 minutes for the inverters to reconnect (the smallest delivering Production Hour is %.1f kWh)"
       % small, all(same_a), same_a)
    ok("ASK A every availability clear of its rounding edge", all(abs(((ga[t]["avail"] * 100) % 1) - 0.5) > 0.001 for t in ga))
    lines25 = [e for e in W.events if e["cause"] == "line" and e["start"].year == 2025]
    f217 = [e for e in Aa.events if e["code"] == "F-217"]
    grp = {}
    for e in f217:
        grp.setdefault(e["raised"].replace(second=0), []).append(e["dev"])
    multi = [k for k, v in grp.items() if len(v) == 16]
    ok("ASK A fairness: each district interruption raised F-217 on all sixteen inverters within the minute",
       len(multi) == len(lines25) == len(Aa.grid) == 4)
    relays25 = [e for e in W.events if e["cause"] == "relay" and e["start"].year == 2025]
    ok("ASK A device A1: every inverter has a relay trip in 2025, each a single-inverter F-217",
       len({e["inv"] for e in relays25}) == 16)
    gap_ok = True
    for (blk, st, hrs) in W.logger_gaps:
        for e in W.events:
            invs = range(16) if e["inv"] is None else [e["inv"]]
            if any(K.BLOCK_OF[j] == blk for j in invs) and not (e["end"] <= st or e["start"] >= st + dt.timedelta(hours=hrs)):
                gap_ok = False
    ok("ASK A device A2: every block logger has gaps in production hours and no gap overlaps an outage",
       gap_ok and len({g[0] for g in W.logger_gaps}) == 4)
    # Ask A's rows are outside the main call's population
    y25 = A.P.year == 2025
    daily_ok, gaps_rows = True, 16 * 8760 - sum(len(v) for v in Aa.hours.values())
    for j, tag in enumerate(AS.TAGS_2025):
        hrs = sorted(Aa.hours[tag])
        reg = Aa.hours[tag][hrs[-1]][1] - Aa.hours[tag][hrs[0]][1] + Aa.hours[tag][hrs[0]][0]
        daily = float(np.nansum(A.P.ac[y25, j]))
        if abs(reg / daily - 1.0) > 0.0005:
            daily_ok = False
    ok("SEPARATION the daily SCADA holds every 2025 inverter-day whole (its AC sums match the registers across the "
       "logger gaps), so the %d missing hourly rows sit only in the ask's file" % gaps_rows,
       daily_ok and gaps_rows > 0 and not np.isnan(A.P.ac[y25]).any())

    # ------------------------------------------------------------------ Ask B
    Bb, gb, sb = G["askB"]
    for k, st in sb.items():
        miss = sum(1 for y in gb if AS.rnd(st[y]) != AS.rnd(gb[y]))
        need = {"S1": 9, "S2": 9, "S3": 2, "S4": 2, "S5": 2, "S6": 4, "S7": 1, "S8": 1}[k.split()[0]]
        ok("ASK B stop %s misses %d of 10" % (k, miss), miss >= need)
    ok("ASK B every year clear of its rounding edge by more than the data's 0.1 kWh resolution",
       all(abs((abs(gb[y]) % 1) - 0.5) > 0.001 for y in gb), min(abs((abs(gb[y]) % 1) - 0.5) for y in gb))
    leg = {(sp, ch, d): k for sp, ch, d, k, end in Bb.legacy}
    mdm1 = {(sp, ch, d): k for sp, ch, d, k, rt, v in Bb.mdm if (sp, ch, d) in leg}
    ok("ASK B fairness: the legacy file and the meter data system agree on every overlapping day",
       len(mdm1) > 400 and all(abs(leg[x] - mdm1[x]) < 0.05 for x in mdm1))
    ok("ASK B device B1: service point 0418 carries DEL rows only until April 17, 2024",
       max(d for sp, ch, d, k, rt, v in Bb.mdm if sp == "0418") < K.STATION_SERVICE_END)
    ok("ASK B device B2: the confirmations amend 2021 and 2024", len(Bb.amend) == 2)
    ok("ASK B the policy makes the confirmation, as amended, the record of a trade's quantity (5.2); the blotter "
       "as booked is stop S3", "is the record of a trade's quantity" in pol and
       any(k.startswith("S3 ") for k in sb))
    # the cracker: landed the call, ran the battery (dedup versions, the overlap), executed every filed rule
    crackA = sa["S4 every F-217 hour excluded, gaps read through the register"]
    crackB = sb["S2 the whole account netted"]
    hitA = sum(1 for t in ga if round(crackA[t]["avail"], 2) == round(ga[t]["avail"], 2))
    hitB = sum(1 for y in gb if AS.rnd(crackB[y]) == AS.rnd(gb[y]))
    ok("CRACKER keeps at most 2 of 26 ask figures", hitA + hitB <= 2, "%d + %d" % (hitA, hitB))

    # ------------------------------------------------------------------ realism sweep
    tags_old = set(K.OLD_TAGS)
    yearly = {}
    for j in range(16):
        for y in range(2016, 2026):
            yearly.setdefault(y, []).append(round(float(np.nansum(A.P.dc[A.P.year == y, j])), 0))
    ok("REALISM no two inverters share a year's DC energy", all(len(set(v)) == 16 for v in yearly.values()))
    late = []
    for f, t in texts.items():
        for m in re.finditer(r"\b(20\d\d)-(\d\d)-(\d\d)\b", t):
            if m.group(0) > "2026-10-06":
                late.append((f, m.group(0)))
    ok("REALISM no document dated after the pull (October 6, 2026)", not late, late[:3])
    names = set(K.PEOPLE.values())
    found = set(re.findall(r"\b([A-Z][a-z]+ [A-Z][a-z]+)\b", alltext))
    persons = {n for n in found if n.split()[0] in {x.split()[0] for x in names}}
    ok("REALISM every person named in the documents is a drawn persona", persons <= names, sorted(persons - names))
    em = [f for f in os.listdir(target) if f.endswith((".csv", ".md", ".txt")) and "\u2014" in open(os.path.join(target, f), encoding="utf-8").read()]
    em += [f for f, t in texts.items() if "\u2014" in t or "\u2013" in t]
    ok("REALISM no em or en dash in the pack", not em, em)
    sub = open(os.path.join(os.path.dirname(os.path.abspath(target)), "submission.md"), encoding="utf-8").read()
    ok("REALISM no em dash in the submission", "\u2014" not in sub and "\u2013" not in sub)

    readme = texts[DOC.F_README]
    unlisted = [f for f in sorted(os.listdir(target)) if f != DOC.F_README and f not in readme]
    ok("REGISTER the README names every shipped file (H9)", not unlisted, unlisted)
    sub_files = set(re.findall(r"\b[\w.-]+\.(?:csv|xlsx|pdf|txt|docx|md)\b", sub)) - {DL.GOLD_BOOK, DL.GOLD_NOTE}
    ok("REGISTER every file the submission names ships (H8)", sub_files <= set(os.listdir(target)), sorted(sub_files - set(os.listdir(target))))

    # ------------------------------------------------------------------ input gates and distractors
    files = sorted(os.listdir(target))
    fmts = {f.rsplit(".", 1)[1] for f in files}
    big = max(B.counts.values())
    ok("GATE 10 or more files, 3 or more formats, one of 25,000 or more rows", len(files) >= 10 and len(fmts) >= 3 and big >= 25000,
       "%d files, %s, %d rows" % (len(files), sorted(fmts), big))
    meta = json.load(open(meta_path))
    ok("GATE two distractors named in metadata.json and shipped", len(meta["distractor_files"]) >= 2 and
       all(d in files for d in meta["distractor_files"]))
    tmp = tempfile.mkdtemp(dir=os.path.dirname(os.path.abspath(target)), suffix=".tmp")
    try:
        for f in files:
            if f not in meta["distractor_files"]:
                shutil.copy(os.path.join(target, f), os.path.join(tmp, f))
        A4 = CA.analyse(tmp, CA.report_gain(tmp))
        _, ga4, _ = AS.ask_a(tmp)
        _, gb4, _ = AS.ask_b(tmp)
        ok("GATE the golden and both asks unchanged with the distractors deleted", A4.volume == vol and
           abs(A4.P90 - A.P90) < 1e-6 and all(ga4[t]["avail"] == ga[t]["avail"] for t in ga) and all(gb4[y] == gb[y] for y in gb))
    finally:
        shutil.rmtree(tmp)

    # ------------------------------------------------------------------ goldens agree with figures.json
    F = json.load(open(os.path.join(os.path.dirname(os.path.abspath(target)), "generator", "figures.json")))
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(golden, DL.GOLD_BOOK), data_only=True)
    ws = wb["Volume 2027"]
    ok("GOLDEN the workbook's volume, P50 and SD are the figures'", ws["B4"].value == F["call"]["volume_mwh"] and
       ws["B5"].value == F["components"]["p50_mwh"] and ws["B6"].value == F["components"]["sd_mwh"])
    wa = wb["Inverter availability 2025"]
    ok("GOLDEN the workbook's availabilities are the figures'", all("%.2f" % wa.cell(row=i, column=5).value ==
       F["ask_a_availability_pct"][wa.cell(row=i, column=1).value] for i in range(2, 18)))
    wh = wb["Hedge record 2016-2025"]
    ok("GOLDEN the workbook's over and under are the figures'", all(wh.cell(row=i, column=4).value ==
       F["ask_b_over_under_mwh"][str(wh.cell(row=i, column=1).value)] for i in range(2, 12)))
    with pdfplumber.open(os.path.join(golden, DL.GOLD_NOTE)) as pdf:
        nt = "\n".join(pg.extract_text() or "" for pg in pdf.pages)
        npages = len(pdf.pages)
    ok("GOLDEN the note is one page and carries the volume, P50 and SD", npages == 1 and "70,000 MWh" in nt and
       "73,000 MWh" in nt and "%s MWh" % format(F["components"]["sd_mwh"], ",") in nt)

    fails = [x for x in RESULTS if not x[1]]
    for label, good, detail in RESULTS:
        print("%s  %s%s" % ("PASS" if good else "FAIL", label, ("  [%s]" % detail) if detail else ""))
    print("checks: %d assertions, %d failed" % (len(RESULTS), len(fails)))
    if fails:
        raise SystemExit("checks failed: %s" % "; ".join(x[0] for x in fails))
    return RESULTS


def _no_repeats(W):
    """No 7-day run of an inverter's daily DC energy appears twice anywhere in the record."""
    seen = set()
    for j in range(K.N_INV):
        v = np.round(W.dc_d[:, j], 1)
        for i in range(0, len(v) - 7):
            key = tuple(v[i:i + 7])
            if all(x == 0 for x in key):
                continue
            if key in seen:
                return False
            seen.add(key)
    return True
