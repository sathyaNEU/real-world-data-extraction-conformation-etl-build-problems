"""The goldens: hedge_volume_2027.xlsx (the volume, the P50 and standard deviation, the performance build and the 25
weather years, each inverter's 2025 availability, each year's over or under) and hedge_note_2027.pdf (the one-page
note with the weather-year chart), then generator/figures.json and submission.md. Every figure is computed from
target/ through calc.py and asks.py."""
from __future__ import annotations

import io
import json
import math
import os

import numpy as np

import asks as AS
import calc as CA
import common as K
import postprocess as PP

GOLD_BOOK = "hedge_volume_2027.xlsx"
GOLD_NOTE = "hedge_note_2027.pdf"


def r_nearest(x, step):
    return int(math.floor(x / step + 0.5)) * step


def compute(target):
    G = {}
    gain = CA.report_gain(target)
    A = CA.analyse(target, gain)
    G["A"] = A
    G["gain"] = gain
    G["gain_pct"] = round(100.0 * (gain - 1.0), 1)
    G["PR_B"] = A.PR_B
    G["wear"] = A.wear
    G["wear_pct"] = 100.0 * (1.0 - A.wear)
    G["PR_new"] = A.PR_new
    G["P50"], G["SD"], G["P90"], G["volume"] = A.P50, A.SD, A.P90, A.volume
    G["P50_r"] = r_nearest(A.P50, 1000)
    G["SD_r"] = r_nearest(A.SD, 100)
    G["E"] = {y: float(e) for y, e in zip(range(2001, 2026), A.E)}
    G["H"] = {y: float(h) for y, h in zip(range(2001, 2026), A.H)}
    G["routes"] = CA.routes(A)
    G["askA"] = AS.ask_a(target)
    G["askB"] = AS.ask_b(target)
    return G


def _label_tag(t):
    return t


def write_book(path, G):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Volume 2027"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E8E1D3")
    rows = [
        ("Hadley Ranch Solar: 2027 forward sale at the one-year P90", None, None),
        ("Energy and commodities desk, for the treasurer's approval, October 2026", None, None),
        (None, None, None),
        ("Volume to sell forward for 2027 (MWh)", G["volume"], "one-year P90 rounded down to the nearest 1,000 MWh (policy 4.1)"),
        ("P50 (MWh)", G["P50_r"], "mean of the 25 weather-year outputs, to the nearest 1,000 MWh"),
        ("Standard deviation (MWh)", G["SD_r"], "sample standard deviation of the 25 weather-year outputs, to the nearest 100 MWh"),
        (None, None, None),
        ("Performance build", None, None),
        ("Performance ratio, October 2025 to September 2026", round(100.0 * G["PR_B"], 1),
         "per cent: net export at the revenue meter over plane-of-array irradiation times 40,320 kWp"),
        ("Original modules' own wear since 2016", int(K.round_half_up(G["wear_pct"], 0)),
         "per cent: fall in DC energy per kWp per kWh/m2 at the inverters' DC input, temperature-corrected, on the days "
         "after each block's wash"),
        ("Side-by-side gain, CV-450BHC over unexposed KS-315P", G["gain_pct"],
         "per cent: twelve months' DC energy per kWp, Corvane report FA-2026-114"),
        ("Performance ratio of the repowered plant", int(K.round_half_up(100.0 * G["PR_new"], 0)),
         "per cent: base-period ratio restored by the wear alone, times the gain"),
        ("Repowered DC capacity (kWp)", 37324.8, "82,944 CV-450BHC modules of 450 W"),
        (None, None, None),
        ("Losses the plant keeps in 2027: the inverters' conversion losses and outages, and the dust between December "
         "washes under the current O&M agreement, all as in the base period.", None, None),
    ]
    for i, (a, b, c) in enumerate(rows, start=1):
        ws.cell(row=i, column=1, value=a)
        if b is not None:
            ws.cell(row=i, column=2, value=b)
        if c is not None:
            ws.cell(row=i, column=3, value=c)
        if i in (1, 8):
            ws.cell(row=i, column=1).font = Font(bold=True, size=12 if i == 1 else 11)
        if i in (4, 5, 6):
            ws.cell(row=i, column=1).font = bold
            ws.cell(row=i, column=2).font = bold
            ws.cell(row=i, column=2).number_format = "#,##0"
    ws.cell(row=9, column=2).number_format = "0.0"
    ws.cell(row=11, column=2).number_format = "0.0"
    ws.cell(row=13, column=2).number_format = "#,##0.0"
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 13
    ws.column_dimensions["C"].width = 96
    for i in (15,):
        ws.cell(row=i, column=1).alignment = Alignment(wrap_text=False)
    # weather years
    w = wb.create_sheet("Weather years")
    for j, h in enumerate(["Weather year", "Plane-of-array irradiation (kWh/m2)", "2027 output (MWh)"], start=1):
        c = w.cell(row=1, column=j, value=h)
        c.font = bold
        c.fill = head_fill
    for i, y in enumerate(range(2001, 2026), start=2):
        w.cell(row=i, column=1, value=y)
        w.cell(row=i, column=2, value=round(G["H"][y], 1)).number_format = "#,##0.0"
        w.cell(row=i, column=3, value=r_nearest(G["E"][y], 100)).number_format = "#,##0"
    n = 27
    w.cell(row=n, column=1, value="P50").font = bold
    w.cell(row=n, column=3, value=G["P50_r"]).number_format = "#,##0"
    w.cell(row=n + 1, column=1, value="Standard deviation").font = bold
    w.cell(row=n + 1, column=3, value=G["SD_r"]).number_format = "#,##0"
    w.cell(row=n + 2, column=1, value="Volume (P90, rounded down)").font = bold
    w.cell(row=n + 2, column=3, value=G["volume"]).number_format = "#,##0"
    w.column_dimensions["A"].width = 26
    w.column_dimensions["B"].width = 34
    w.column_dimensions["C"].width = 18
    # availability
    A, gold, stops = G["askA"]
    a = wb.create_sheet("Inverter availability 2025")
    for j, h in enumerate(["Inverter", "Production hours", "Unavailable hours", "Excluded hours", "Availability (%)"], start=1):
        c = a.cell(row=1, column=j, value=h)
        c.font = bold
        c.fill = head_fill
    for i, tag in enumerate(AS.TAGS_2025, start=2):
        v = gold[tag]
        a.cell(row=i, column=1, value=tag)
        a.cell(row=i, column=2, value=v["P"])
        a.cell(row=i, column=3, value=v["U"])
        a.cell(row=i, column=4, value=v["X"])
        a.cell(row=i, column=5, value=round(v["avail"], 2)).number_format = "0.00"
    a.cell(row=19, column=1, value="Schedule 4 of the O&M agreement: Excluded Hours are the district's 69 kV interruptions "
           "and the stops the owner requested.")
    for c, wd in zip("ABCDE", (12, 17, 18, 15, 16)):
        a.column_dimensions[c].width = wd
    # hedge record
    B, gb, sb = G["askB"]
    net = AS.net_by_year(B)
    sold = AS.sold_by_year(B)
    h = wb.create_sheet("Hedge record 2016-2025")
    for j, t in enumerate(["Year", "Net energy at the delivery point (MWh)", "Volume sold forward (MWh)",
                           "Over (+) or under (-) (MWh)"], start=1):
        c = h.cell(row=1, column=j, value=t)
        c.font = bold
        c.fill = head_fill
    for i, y in enumerate(range(2016, 2026), start=2):
        h.cell(row=i, column=1, value=y)
        h.cell(row=i, column=2, value=AS.rnd(net[y])).number_format = "#,##0"
        h.cell(row=i, column=3, value=AS.rnd(sold[y])).number_format = "#,##0"
        h.cell(row=i, column=4, value=AS.rnd(gb[y])).number_format = "#,##0;-#,##0"
    h.cell(row=13, column=1, value="Delivery point: district service point 0417 (REC less DEL). Volumes net of the 2019 "
           "and 2022 buy-backs, with the 2021 and 2024 amendments.")
    for c, wd in zip("ABCD", (8, 36, 26, 26)):
        h.column_dimensions[c].width = wd
    for sh in (w, a, h):
        sh.freeze_panes = "A2"
    n = wb.create_sheet("Notes")
    notes = [
        ("Sources", "Plant SCADA daily exports (original system to March 2, 2020; replacement system from March 3, 2020), "
                    "pulled October 6, 2026; Tresco work orders; Corvane report FA-2026-114; Clearfield site-adapted "
                    "hourly series 2001 to 2025 (delivered January 2026); module replacement scope, July 2026."),
        ("Base period", "October 1, 2025 to September 30, 2026, the twelve complete months before the sale (energy risk "
                        "policy, section 4.4)."),
        ("Performance ratio", "Net export at the revenue meter (SCADA PLANT-RM export less import) over plane-of-array "
                              "irradiation times DC nameplate."),
        ("Wear", "Days 1 to 3 after each block's wash in the CLN work orders, inverters with full operating hours; DC "
                 "energy at the inverter input per kWp per kWh/m2, corrected to 25 degC at -0.42 per cent per degC on the "
                 "irradiance-weighted module temperature (KS-315P datasheet). 2016 against the December 2025 wash."),
        ("Weather years", "Full calendar years 2001 to 2025 of the long-term record, Arizona time (policy section 4.2)."),
        ("Rounding", "Volume rounded down to the nearest 1,000 MWh from the unrounded P90 (policy section 4.1); P50 to the "
                     "nearest 1,000 MWh, standard deviation to the nearest 100 MWh, weather-year outputs to the nearest "
                     "100 MWh."),
        ("Availability", "O&M agreement Schedule 4, contract year 2025: Production Hours from the met station's hourly "
                         "plane-of-array irradiance; an hour missing from the logger export counts as delivered where the "
                         "inverter's AC register advanced across it."),
        ("Hedge record", "Net energy at district service point 0417 (REC less DEL) from the legacy billing export to "
                         "February 29, 2020 and the meter data system after it, latest version of each day; volumes from "
                         "the blotter, with Westmark's amendments of the 2021 and 2024 confirmations."),
    ]
    for i, (k, v) in enumerate(notes, start=1):
        n.cell(row=i, column=1, value=k).font = bold
        c = n.cell(row=i, column=2, value=v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        n.cell(row=i, column=1).alignment = Alignment(vertical="top")
    n.column_dimensions["A"].width = 18
    n.column_dimensions["B"].width = 110
    ws.cell(row=4, column=2).fill = PatternFill("solid", fgColor="F3E3C3")
    wb.save(path)


def chart_png(G):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    for f in ("/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf", "/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf"):
        font_manager.fontManager.addfont(f)
    plt.rcParams["font.family"] = "Carlito"
    years = list(range(2001, 2026))
    vals = [G["E"][y] / 1000.0 for y in years]
    fig, ax = plt.subplots(figsize=(7.4, 3.6), dpi=150)
    ax.bar(years, vals, color="#c9a46a", width=0.7, zorder=2)
    p50, vol = G["P50"] / 1000.0, G["volume"] / 1000.0
    ax.axhline(p50, color="#1f2329", linewidth=1.3, zorder=3)
    ax.axhline(vol, color="#8a1c1c", linewidth=1.4, linestyle="--", zorder=3)
    ax.text(2025.6, p50, " P50\n %s MWh" % format(G["P50_r"], ","), va="center", ha="left", fontsize=7.4, color="#1f2329")
    ax.text(2025.6, vol, " Volume\n %s MWh" % format(G["volume"], ","), va="center", ha="left", fontsize=7.4,
            color="#8a1c1c")
    ax.set_xlim(2000.4, 2028.6)
    ax.set_ylim(60, 78)
    ax.set_xticks(years)
    ax.set_xticklabels(["%02d" % (y % 100) for y in years], fontsize=7.2)
    ax.set_xlabel("Weather year of the long-term record (2001 to 2025)", fontsize=8)
    ax.set_ylabel("2027 output (GWh)", fontsize=8)
    ax.tick_params(axis="y", labelsize=7.5)
    ax.grid(axis="y", color="#e3e3e3", linewidth=0.6, zorder=0)
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    ax.set_title("The volume sits under the plant's 2027 output in all but the cloudiest weather years\n"
                 "Hadley Ranch Solar with the new modules, output under each year's weather (GWh, axis from 60)",
                 fontsize=9.2)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    buf.seek(0)
    return buf


def write_note(path, G):
    from reportlab import rl_config
    rl_config.invariant = 1
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    import docs as DOC
    ink, grey, accent = DOC.INK, DOC.GREY, DOC.MV
    body = ParagraphStyle("nb", fontName="Carlito", fontSize=10.3, leading=13.6, textColor=ink, spaceAfter=6)
    h1 = ParagraphStyle("nh", fontName="Carlito-Bold", fontSize=14.2, leading=18, textColor=accent, spaceAfter=6)
    cell = ParagraphStyle("nc", fontName="Carlito", fontSize=9.4, leading=11.8, textColor=ink)
    cellb = ParagraphStyle("ncb", fontName="Carlito-Bold", fontSize=9.4, leading=11.8, textColor=grey)
    small = ParagraphStyle("nsm", fontName="Carlito", fontSize=8.0, leading=10.2, textColor=grey, spaceAfter=3)
    head = Table([[Paragraph("To", cellb), Paragraph("%s, Treasurer" % K.PEOPLE["treasurer"], cell)],
                  [Paragraph("From", cellb), Paragraph("%s, Energy and Commodities Manager" % K.PEOPLE["desk"], cell)],
                  [Paragraph("Date", cellb), Paragraph("October 8, 2026", cell)],
                  [Paragraph("Subject", cellb), Paragraph("Hadley Ranch Solar: 2027 forward sale volume for approval "
                                                          "(energy risk policy, section 5.1)", cell)]],
                 colWidths=[0.75 * inch, 6.0 * inch])
    head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                              ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5), ("TOPPADDING", (0, 0), (-1, -1), 1.5),
                              ("LINEBELOW", (0, -1), (-1, -1), 0.6, colors.HexColor("#a3abb5"))]))
    wear = int(K.round_half_up(G["wear_pct"], 0))
    prn = int(K.round_half_up(100.0 * G["PR_new"], 0))
    s = [head, Spacer(1, 10),
         Paragraph("Sell %s MWh of Hadley Ranch's 2027 output forward" % format(G["volume"], ","), h1),
         Paragraph("I recommend we sell %s MWh forward by October 31, the plant's one-year P90 for 2027 with its new "
                   "modules, rounded down. Across the 25 weather years of Clearfield's record the plant's 2027 output "
                   "has a P50 of %s MWh and a standard deviation of %s MWh."
                   % (format(G["volume"], ","), format(G["P50_r"], ","), format(G["SD_r"], ",")), body),
         Paragraph("The volume starts from what the plant delivered in the twelve months to September 30, a "
                   "performance ratio of %.1f per cent, and changes only what the replacement changes. The original "
                   "modules' own wear goes with them: measured at the inverters' DC input on the days after each block's "
                   "wash, where the dust and the inverters drop out, it comes to %d per cent since 2016. The new modules "
                   "then add the %.1f per cent per kWp that Corvane measured against unexposed original modules in block 3 "
                   "over the same twelve months. That puts the repowered plant at a performance ratio of %d per cent on "
                   "37,324.8 kWp." % (100.0 * G["PR_B"], wear, G["gain_pct"], prn), body),
         Paragraph("What the replacement does not touch stays in the number: the inverters, which have run on "
                   "refurbished power stages since 2018 and are out more often than they were, and the dust that builds "
                   "between the single December wash the O&amp;M agreement provides through 2027. Carrying the plant's whole "
                   "decline since 2016 back into 2027 would overstate the sale; carrying today's output without the "
                   "wear would understate it.", body),
         Spacer(1, 2),
         Image(chart_png(G), width=6.9 * inch, height=6.9 * inch * 3.6 / 7.4),
         Spacer(1, 4),
         Paragraph("Base period: October 1, 2025 to September 30, 2026, the twelve complete months before the sale "
                   "(policy section 4.4). P50 and standard deviation over the full calendar years 2001 to 2025 of the "
                   "site-adapted record (section 4.2); the P90 is the P50 less 1.2816 standard deviations, rounded down "
                   "to the nearest 1,000 MWh (section 4.1).", small),
         Paragraph("Workbook: hedge_volume_2027.xlsx, with the build, the weather years, each inverter's 2025 "
                   "availability under Schedule 4 and the 2016 to 2025 hedge record.", small)]
    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.85 * inch, rightMargin=0.85 * inch, topMargin=0.7 * inch,
                            bottomMargin=0.7 * inch, title="Hadley Ranch Solar 2027 forward volume",
                            author=K.PEOPLE["desk"], subject="", creator="", producer="")

    def foot(canvas, d):
        canvas.saveState()
        canvas.setFont("Carlito", 7.6)
        canvas.setFillColor(grey)
        canvas.drawString(0.85 * inch, 0.45 * inch, "Marden-Vail Building Products  |  Treasury, energy and commodities desk  "
                          "|  Confidential")
        canvas.restoreState()
    doc.build(s, onFirstPage=foot, onLaterPages=foot)


def figures(G):
    A, gold, stops = G["askA"]
    B, gb, sb = G["askB"]
    F = {
        "call": {"volume_mwh": G["volume"], "volume": format(G["volume"], ","), "hedge_year": 2027},
        "components": {"p50_mwh": G["P50_r"], "p50": format(G["P50_r"], ","), "sd_mwh": G["SD_r"], "sd": format(G["SD_r"], ","),
                       "base_period_pr_pct": round(100.0 * G["PR_B"], 1), "wear_pct": int(K.round_half_up(G["wear_pct"], 0)),
                       "gain_pct": G["gain_pct"], "repowered_pr_pct": int(K.round_half_up(100.0 * G["PR_new"], 0)),
                       "new_kwp": "37,324.8", "old_kwp": "40,320", "old_kwp_inverter": "2,520", "z": 1.2816,
                       "gamma_pct_per_degc": -0.42, "rounding_step": "1,000",
                       "weather_years": 25, "clean_days": 3},
        "unrounded": {"p50": round(G["P50"], 1), "sd": round(G["SD"], 1), "p90": round(G["P90"], 1),
                      "wear_ratio": round(G["wear"], 4), "gain_ratio": round(G["gain"], 4), "pr_base": round(G["PR_B"], 4),
                      "pr_new": round(G["PR_new"], 4)},
        "rivals_p90_mwh": {k: round(v, 0) for k, v in G["routes"].items()},
        "rival_volumes": {"last twelve months with the gain": format(int(G["routes"]["R0 the last twelve months carried with the gain"] // 1000) * 1000, ","),
                          "first year with the gain": format(int(G["routes"]["R1 the first year carried with the gain"] // 1000) * 1000, ","),
                          "whole DC-side decline": format(int(G["routes"]["R2 the DC-side decline on available days"] // 1000) * 1000, ",")},
        "ask_a_availability_pct": {t: "%.2f" % gold[t]["avail"] for t in AS.TAGS_2025},
        "ask_a_hours": {t: {"production": gold[t]["P"], "unavailable": gold[t]["U"], "excluded": gold[t]["X"]} for t in AS.TAGS_2025},
        "ask_b_over_under_mwh": {str(y): AS.rnd(gb[y]) for y in range(2016, 2026)},
        "ask_b_over_under": {str(y): format(AS.rnd(gb[y]), ",") for y in range(2016, 2026)},
        "years": list(range(2016, 2026)) + [2001, 2025, 2026, 2027],
        "references": {"policy_sections": ["4.1", "4.2", "4.4", "3.1"], "schedule": "4", "temperature_c": 25,
                       "service_point": "0417", "days_after_wash": [1, 2, 3], "months": 12},
    }
    return F


def write_submission(path, G):
    A, gold, stops = G["askA"]
    B, gb, sb = G["askB"]
    vol = format(G["volume"], ",")
    L = ["## Tags", "",
         "**Domain:** Energy Systems & Utilities (asset condition and maintenance: a solar plant's output after its "
         "modules are replaced).",
         "**Analytical objective:** Forecasting & Predictive Modeling.", "",
         "## 1. Final Recommendation", "",
         "**Sell %s MWh of Hadley Ranch Solar's 2027 output forward: the one-year P90 of the plant with its new "
         "modules, rounded down to the nearest 1,000 MWh.**" % vol,
         "The P50 is %s MWh with a standard deviation of %s MWh over the 25 weather years. Not the last twelve months "
         "carried onto the new modules with the side-by-side gain, which keeps the original modules' wear that the "
         "replacement removes; not the plant's first year carried with the gain, nor its whole DC-side decline "
         "restored, which bring back the inverters' losses or the dust carried between washes that the plant keeps."
         % (format(G["P50_r"], ","), format(G["SD_r"], ",")), "",
         "## 2. Critical Components", "",
         "1. The original modules' own wear since 2016: the DC energy per kWp per unit of plane-of-array irradiation at "
         "the inverters' DC input, temperature-corrected, on the days after each block's wash, is **%d%%** lower after "
         "the December 2025 wash than in 2016." % int(K.round_half_up(G["wear_pct"], 0)),
         "2. The plant's performance ratio over the base period, October 2025 to September 2026: **%.1f%%**."
         % (100.0 * G["PR_B"]),
         "3. The repowered plant's performance ratio, the base period's restored by the wear alone and raised by the "
         "%.1f%% side-by-side gain: **%d%%**." % (G["gain_pct"], int(K.round_half_up(100.0 * G["PR_new"], 0))),
         "4. P50 **%s MWh** and standard deviation **%s MWh**." % (format(G["P50_r"], ","), format(G["SD_r"], ",")), "",
         "## 3. Step-by-Step Solution", "",
         "1. Joined hrs_scada_daily_report_2015-11-20_to_2020-03-02.csv (local days, kWh) and "
         "hadley_scada_daily_2020-03-03_to_2026-09-30.csv (UTC day starts, MWh) through scada_tag_crosswalk.csv into "
         "one daily record per inverter, B2-INV3R continuing inverter 07.",
         "2. Base period (policy section 4.4), October 1, 2025 to September 30, 2026: PLANT-RM E_EXP less E_IMP over "
         "H_POA times 40,320 kWp gives a performance ratio of %.1f%%." % (100.0 * G["PR_B"]),
         "3. Took each block's washes from the CLN work orders in tresco_work_orders_hadley_ranch_2015-2026.xlsx and, on "
         "the inverter-days 1 to 3 days after each wash with full operating hours, E_DC over H_POA times 2,520 kWp "
         "corrected to 25 degC with the KS-315P's -0.42 %%/degC on T_MOD_IW: the level after the December 2025 wash is "
         "%d%% below 2016's, the original modules' own wear." % int(K.round_half_up(G["wear_pct"], 0)),
         "4. On the same record the inverters' AC-to-DC ratio has fallen and the DC energy on every available day has "
         "fallen further than on the days after the washes, now one a year under Schedule 2 of the O&M agreement: "
         "losses the replacement scope keeps, so the base period carries them into 2027.",
         "5. Divided the base period's performance ratio by the wear and multiplied by the %.1f%% gain of "
         "corvane_fa-2026-114_hadley_ranch_side_by_side.pdf: %d%% on 37,324.8 kWp."
         % (G["gain_pct"], int(K.round_half_up(100.0 * G["PR_new"], 0))),
         "6. Applied it to each full year's plane-of-array irradiation in clearfield_hadley_ranch_poa_hourly_2001-2025.csv "
         "(local calendar years 2001 to 2025): P50 %s MWh, sample standard deviation %s MWh (policy section 4.2)."
         % (format(G["P50_r"], ","), format(G["SD_r"], ",")),
         "7. P50 less 1.2816 standard deviations, rounded down to the nearest 1,000 MWh (policy section 4.1): %s MWh."
         % vol, "",
         "## 4. Deliverable Answers", "",
         "### hedge_volume_2027.xlsx", "",
         "1. Volume to sell forward for 2027: %s MWh." % vol,
         "2. P50: %s MWh." % format(G["P50_r"], ","),
         "3. Standard deviation: %s MWh." % format(G["SD_r"], ","),
         "4. Each inverter's 2025 availability under Schedule 4 of the O&M agreement, per cent:"]
    for t in AS.TAGS_2025:
        L.append("   - %s: %.2f%%" % (t, gold[t]["avail"]))
    L.append("5. Net energy at the delivery point over (+) or under (-) the volume sold forward, MWh:")
    for y in range(2016, 2026):
        v = AS.rnd(gb[y])
        L.append("   - %d: %s%s MWh" % (y, "+" if v > 0 else "", format(v, ",")))
    L += ["", "### hedge_note_2027.pdf", "",
          "1. One-page note to the treasurer recommending the %s MWh sale, with a bar chart of the repowered "
          "plant's 2027 output in each weather year 2001 to 2025 and lines marking the P50 (%s MWh) and the volume." % (
              vol, format(G["P50_r"], ",")), ""]
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L))


def main(target, golden, figures_path, submission_path, B=None):
    os.makedirs(golden, exist_ok=True)
    G = compute(target)
    write_book(os.path.join(golden, GOLD_BOOK), G)
    PP.normalize_ooxml(os.path.join(golden, GOLD_BOOK), K.PEOPLE["desk"], K.PEOPLE["desk"], "2026-10-07T16:05:00Z",
                       "2026-10-08T22:31:00Z", excel=True, company=K.COMPANY)
    write_note(os.path.join(golden, GOLD_NOTE), G)
    PP.normalize_pdf(os.path.join(golden, GOLD_NOTE), b"D:20261008153507-07'00'", b"D:20261008153507-07'00'")
    F = figures(G)
    write_submission(submission_path, G)
    with open(figures_path, "w") as fh:
        json.dump(F, fh, indent=1)
        fh.write("\n")
    return G
