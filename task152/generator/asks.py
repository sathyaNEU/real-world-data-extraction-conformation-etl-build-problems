"""The two ask tables, computed from the shipped files: each inverter's 2025 availability as the O&M agreement's
Schedule 4 defines it, and each year's over or under of the plant's net energy at the delivery point against the
volume sold forward. Each wrong pipeline the ask sheet names (its stops) is computed here on the same reads."""
from __future__ import annotations

import csv
import datetime as dt
import math
import os
import re

import pdfplumber
from openpyxl import load_workbook

import common as K
import docs as DOC
import records as R

TAGS_2025 = [K.NEW_TAG_07R if i == K.INV07 else K.NEW_TAGS[i] for i in range(K.N_INV)]


def _utc(s):
    return dt.datetime.strptime(s, "%Y-%m-%dT%H:%M:%SZ")


# ---------------------------------------------------------------------------------------------------- Ask A
class AskA:
    pass


def read_ask_a(folder, align="end"):
    """align: how the hourly AC file's stamps meet the met station's local hour starts. 'end' (each row is the hour
    ending at its UTC stamp), 'begin' (the stamp read as the hour's start), 'utc' (the UTC stamp read as local)."""
    A = AskA()
    A.prod = set()
    with open(os.path.join(folder, R.F_MET), encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            if float(row["poa_wm2"]) >= 50.0:
                A.prod.add(dt.datetime.strptime(row["hour_beginning_mst"], "%Y-%m-%d %H:%M"))
    A.hours = {t: {} for t in TAGS_2025}      # local hour start -> (ac_kwh, register at the hour's end)
    with open(os.path.join(folder, R.F_HOURLY), encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            end = _utc(row["interval_end_utc"])
            start = {"end": end - R.UTC_OFF - dt.timedelta(hours=1), "begin": end - R.UTC_OFF,
                     "utc": end - dt.timedelta(hours=1)}[align]
            A.hours[row["device"]][start] = (float(row["ac_kwh"]), int(row["ac_register_kwh"]))
    A.events = []
    with open(os.path.join(folder, R.F_EVENTS), encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            A.events.append(dict(dev=row["device"], code=row["code"], raised=_utc(row["raised_utc"]),
                                 cleared=_utc(row["cleared_utc"])))
    # the district's interruptions (the report's table)
    A.grid = []
    with pdfplumber.open(os.path.join(folder, DOC.F_INTERRUPT)) as pdf:
        text = "\n".join(pg.extract_text() or "" for pg in pdf.pages)
    for m in re.finditer(r"(HR-25-\d{3}) (\d\d/\d\d/\d{4} \d\d:\d\d) (\d\d/\d\d/\d{4} \d\d:\d\d)", text):
        A.grid.append((dt.datetime.strptime(m.group(2), "%m/%d/%Y %H:%M"), dt.datetime.strptime(m.group(3), "%m/%d/%Y %H:%M")))
    # the owner's stops (the work orders)
    A.owner = []
    wb = load_workbook(os.path.join(folder, R.F_WO), read_only=True)
    first = True
    for row in wb.active.iter_rows(values_only=True):
        if first:
            first = False
            continue
        no, typ, asset, desc, made, done = row[:6]
        if typ != "OWN" or done.year != 2025:
            continue
        m = re.search(r"stop (.+?) for .+?Stopped (\d\d:\d\d), restarted (\d\d:\d\d)", desc)
        devs = [x.strip() for x in m.group(1).split(",")]
        a = dt.datetime.combine(done.date(), dt.datetime.strptime(m.group(2), "%H:%M").time())
        b = dt.datetime.combine(done.date(), dt.datetime.strptime(m.group(3), "%H:%M").time())
        A.owner.append((devs, a, b))
    return A


def _overlaps(t, a, b):
    """The clock hour starting at t overlaps the window [a, b)."""
    return t < b and t + dt.timedelta(hours=1) > a


def availability(A, gaps="register", exclude="grid+owner", local_events=False, derate_unavailable=False):
    """Schedule 4 availability per inverter. gaps: 'register' (an hour missing from the export counts as delivered
    when the register advanced across the gap), 'zero' (counts as no AC energy). exclude: which zero-energy hours are
    Excluded Hours: 'grid+owner' (the district's interruptions and the owner's stops), 'none', 'f217+owner' (every
    hour an F-217 alarm was open, plus the owner's stops), 's020' (the district's interruptions and every hour an S-020
    stop was open). local_events: the event log's UTC stamps read as local time. derate_unavailable: hours with a W-410
    derate open count as unavailable."""
    import bisect
    off = dt.timedelta(0) if local_events else R.UTC_OFF
    out = {}
    prod = sorted(A.prod)
    for tag in TAGS_2025:
        rows = A.hours[tag]
        have = sorted(rows)
        P = U = X = 0
        f217 = [(e["raised"] - off, e["cleared"] - off) for e in A.events if e["dev"] == tag and e["code"] == "F-217"]
        s020 = [(e["raised"] - off, e["cleared"] - off) for e in A.events if e["dev"] == tag and e["code"] == "S-020"]
        w410 = [(e["raised"] - off, e["cleared"] - off) for e in A.events if e["dev"] == tag and e["code"] == "W-410"]
        own = [(a, b) for devs, a, b in A.owner if tag in devs]
        for t in prod:
            if t in rows:
                delivered = rows[t][0] > 0.0
            elif gaps == "zero":
                delivered = False
            else:
                k = bisect.bisect_left(have, t)
                if k == 0 or k == len(have):
                    delivered = False
                else:
                    lo = rows[have[k - 1]][1]
                    hi = rows[have[k]][1] - rows[have[k]][0]
                    delivered = hi - lo > 0.5
            P += 1
            if derate_unavailable and any(_overlaps(t, a, b) for a, b in w410):
                U += 1
                continue
            if delivered:
                continue
            if exclude == "none":
                excl = False
            elif exclude == "grid+owner":
                excl = any(_overlaps(t, a, b + dt.timedelta(minutes=10)) for a, b in A.grid) or \
                    any(_overlaps(t, a, b) for a, b in own)
            elif exclude == "f217+owner":
                excl = any(_overlaps(t, a, b) for a, b in f217) or any(_overlaps(t, a, b) for a, b in own)
            elif exclude == "s020":
                excl = any(_overlaps(t, a, b + dt.timedelta(minutes=10)) for a, b in A.grid) or \
                    any(_overlaps(t, a, b) for a, b in s020)
            else:
                raise ValueError(exclude)
            if excl:
                X += 1
            else:
                U += 1
        out[tag] = dict(P=P, U=U, X=X, avail=100.0 * (1.0 - U / (P - X)))
    return out


def ask_a(folder):
    A = read_ask_a(folder)
    gold = availability(A)
    stops = {
        "S1 every zero or missing hour unavailable, nothing excluded": availability(A, gaps="zero", exclude="none"),
        "S2 every F-217 hour excluded, missing hours as zero": availability(A, gaps="zero", exclude="f217+owner"),
        "S3 the district's outages excluded, missing hours as zero": availability(A, gaps="zero"),
        "S4 every F-217 hour excluded, gaps read through the register": availability(A, exclude="f217+owner"),
        "S5 the event log's stamps read as local time": availability(A, exclude="f217+owner", local_events=True),
        "S6 every S-020 stop excluded": availability(A, exclude="s020"),
        "S7 derate hours counted unavailable": availability(A, derate_unavailable=True),
        "S8 the hourly file's hour-ending stamps read as hour-beginning": availability(read_ask_a(folder, "begin")),
        "S9 the hourly file's UTC stamps read as Arizona time": availability(read_ask_a(folder, "utc")),
    }
    return A, gold, stops


# ---------------------------------------------------------------------------------------------------- Ask B
class AskB:
    pass


def read_ask_b(folder):
    B = AskB()
    B.legacy = []         # (sp, channel, day, kwh, stamp date)
    with open(os.path.join(folder, R.F_LEGACY), encoding="utf-8") as fh:
        for line in fh:
            line = line.rstrip("\r\n")
            if not line:
                continue
            sp = line[10:16].strip()
            ch = line[26:30].strip()
            end = dt.datetime.strptime(line[30:43], "%Y%m%d %H%M")
            kwh = float(line[43:57])
            day = (end - dt.timedelta(minutes=1)).date()
            B.legacy.append((sp, ch, day, kwh, end.date()))
    B.mdm = []
    with open(os.path.join(folder, R.F_MDM), encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            B.mdm.append((row["service_point"], row["channel"], dt.date.fromisoformat(row["read_date"]),
                          float(row["kwh"]), row["read_type"], int(row["version"])))
    wb = load_workbook(os.path.join(folder, R.F_BLOTTER), read_only=True)
    B.trades = []
    first = True
    for row in wb.active.iter_rows(values_only=True):
        if first:
            first = False
            continue
        tid, tdate, cpty, prod, deliv, side, qty, price = row[:8]
        B.trades.append(dict(id=tid, year=int(str(deliv)[2:]), side=side, qty=float(qty)))
    with pdfplumber.open(os.path.join(folder, DOC.F_CONF)) as pdf:
        text = " ".join((pg.extract_text() or "").replace("\n", " ") for pg in pdf.pages)
    B.amend = {}
    for m in re.finditer(r"reference (MVB-\d{4}-\d{3})\. The Contract Quantity is amended to ([\d,]+) MWh", text):
        B.amend[m.group(1)] = float(m.group(2).replace(",", ""))
    return B


def net_by_year(B, sps=("0417",), stack=False, versions="latest", legacy_year="day"):
    """Each year's REC less DEL over the service points given, in MWh: the legacy file to February 29, 2020 and the
    meter data system after it (stack: both files summed where they overlap), the meter data system's latest version
    of each day (versions 'all': every version summed; 'first': each day's first), the legacy totals dated by their day (legacy_year 'stamp':
    by the date in the interval-end stamp)."""
    recs = []
    legacy_days = set()
    for sp, ch, d, k, end in B.legacy:
        recs.append((sp, ch, end.year if legacy_year == "stamp" else d.year, k))
        legacy_days.add((sp, ch, d))
    mdm = {}
    for sp, ch, d, k, rt, v in B.mdm:
        mdm.setdefault((sp, ch, d), []).append((v, k))
    for key in sorted(mdm):
        if key in legacy_days and not stack:
            continue
        vs = sorted(mdm[key])
        for v, k in (vs if versions == "all" else vs[:1] if versions == "first" else vs[-1:]):
            recs.append((key[0], key[1], key[2].year, k))
    tot = {}
    for sp, ch, y, k in recs:
        if sp in sps:
            tot[y] = tot.get(y, 0.0) + (k if ch == "REC" else -k)
    return {y: v / 1000.0 for y, v in tot.items()}


def sold_by_year(B, amendments=True, sides=True):
    out = {}
    for t in B.trades:
        q = B.amend.get(t["id"], t["qty"]) if amendments else t["qty"]
        sgn = (1.0 if t["side"] == "S" else -1.0) if sides else 1.0
        out[t["year"]] = out.get(t["year"], 0.0) + sgn * q
    return out


def over_under(net, sold):
    return {y: net[y] - sold[y] for y in range(2016, 2026)}


def ask_b(folder):
    B = read_ask_b(folder)
    gold = over_under(net_by_year(B), sold_by_year(B))
    stops = {
        "S1 the whole account netted, the blotter as booked": over_under(net_by_year(B, sps=("0417", "0418")), sold_by_year(B, amendments=False)),
        "S2 the whole account netted": over_under(net_by_year(B, sps=("0417", "0418")), sold_by_year(B)),
        "S3 the blotter as booked (amendments missed)": over_under(net_by_year(B), sold_by_year(B, amendments=False)),
        "S4 both meter files stacked": over_under(net_by_year(B, stack=True), sold_by_year(B)),
        "S5 the buy-backs added": over_under(net_by_year(B), sold_by_year(B, sides=False)),
        "S6 the legacy totals dated by their stamps": over_under(net_by_year(B, legacy_year="stamp"), sold_by_year(B)),
        "S7 every meter data version summed": over_under(net_by_year(B, versions="all"), sold_by_year(B)),
        "S8 the meter data system's first version of each day": over_under(net_by_year(B, versions="first"), sold_by_year(B)),
    }
    return B, gold, stops


def rnd(x):
    return int(math.floor(x + 0.5)) if x >= 0 else -int(math.floor(-x + 0.5))
