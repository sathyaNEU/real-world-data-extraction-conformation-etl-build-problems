"""task197: the remaining pack files, the ask quantities, the goldens and figures.json."""
import os, csv, json, datetime as dt
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.chart import BarChart, Reference, Series
from openpyxl.chart.marker import Marker
import sim, core, docs, build

HV_SITES = {"VKF-ZWT", "VKF-DLK"}        # connections above 110 kV carry a grid loss adjustment
LOSS_ADJ = 1.0042
PRELIM_GAP = ("VKF-NWV", dt.date(2026, 10, 6))   # the day the preliminary October statement missed
CONTROLLER = {"battery": "Nedvolt NX-480 droop module", "gas_engine": "Stedinck GCU-7 governor",
              "industrial_load": "Terborgh LM-220 load controller"}
VOLTS = {"VKF-OOS": 50, "VKF-ZWT": 150, "VKF-HGZ": 20, "VKF-KPV": 50, "VKF-MPD": 20,
         "VKF-NWV": 50, "VKF-TRB": 20, "VKF-DLK": 150, "VKF-RDL": 20}
COMMISSIONED = {"VKF-OOS": "2023-04-18", "VKF-ZWT": "2022-11-02", "VKF-HGZ": "2019-06-25",
                "VKF-KPV": "2024-02-29", "VKF-MPD": "2018-09-12", "VKF-NWV": "2023-08-07",
                "VKF-TRB": "2021-03-15", "VKF-DLK": "2025-01-20", "VKF-RDL": "2017-10-30"}
PLACE = {"VKF-OOS": "Oosterhout, Noord-Brabant", "VKF-ZWT": "Zwartewater, Overijssel",
         "VKF-HGZ": "Hoogezand, Groningen", "VKF-KPV": "Kampvliet, Zeeland",
         "VKF-MPD": "Meppelerdiep, Drenthe", "VKF-NWV": "Nieuwveen, Zuid-Holland",
         "VKF-TRB": "Terborgh, Gelderland", "VKF-DLK": "Dollardkust, Groningen",
         "VKF-RDL": "Reiderland, Groningen"}


# Every event the control platform logged that is not a reserve withdrawal. The archive's capture
# placement reads the same list, so the two files agree on when each event happened.
BENCH = [
        ("2026-09-15T07:42:11Z", "VKF-KPV", "firmware", "FW-APPLY", "inverter droop module firmware 4.18.2 applied"),
        ("2026-09-15T09:05:44Z", "VKF-ZWT", "firmware", "FW-APPLY", "inverter droop module firmware 4.18.2 applied"),
        ("2026-09-17T13:21:08Z", "VKF-TRB", "firmware", "FW-APPLY", "load controller firmware 2.6.0 applied"),
        ("2026-09-18T06:55:19Z", "VKF-HGZ", "firmware", "FW-APPLY", "governor firmware 11.4 applied to engines 1 to 4"),
        ("2026-09-21T05:02:37Z", "VKF-NWV", "mode", "MOD-AUTO", "platform control returned to automatic"),
        ("2026-09-22T22:13:50Z", "VKF-DLK", "alarm", "ALM-TEMP", "inverter bank 1 temperature alarm, cleared without withdrawal"),
        ("2026-09-24T04:48:02Z", "VKF-KPV", "mode", "MOD-LOCAL", "site on local control for planned works, out of reserve market"),
        ("2026-09-25T23:59:12Z", "VKF-KPV", "mode", "MOD-AUTO", "platform control returned to automatic"),
        ("2026-09-26T11:07:31Z", "VKF-OOS", "setpoint", "SET-OVR", "manual programme override, 2 periods"),
        ("2026-09-29T03:31:44Z", "VKF-DLK", "mode", "MOD-LOCAL", "site on local control for protection test, out of reserve market"),
        ("2026-09-29T21:44:55Z", "VKF-DLK", "mode", "MOD-AUTO", "platform control returned to automatic"),
        ("2026-10-01T08:12:26Z", "VKF-OOS", "metering", "MET-CHG", "metering point replaced at connection, new EAN effective 2026-10-01"),
        ("2026-10-02T15:38:09Z", "VKF-RDL", "alarm", "ALM-GAS", "gas pressure warning, cleared without withdrawal"),
        ("2026-10-04T02:16:40Z", "VKF-HGZ", "archive", "ARC-RST", "high resolution archive restored from standby after storage fault"),
        ("2026-10-06T04:05:13Z", "VKF-MPD", "mode", "MOD-LOCAL", "engine overhaul, out of reserve market"),
        ("2026-10-06T23:51:02Z", "VKF-MPD", "mode", "MOD-AUTO", "platform control returned to automatic"),
        ("2026-10-08T12:29:58Z", "VKF-ZWT", "setpoint", "SET-OVR", "manual programme override, 1 period"),
        ("2026-10-09T17:02:44Z", "VKF-TRB", "alarm", "ALM-COM", "communications warning, cleared without withdrawal"),
        ("2026-10-11T19:47:36Z", "VKF-KPV", "alarm", "ALM-TEMP", "cell string temperature alarm, cleared without withdrawal"),
]


def event_log(path, placed):
    rows, eid = [], 4100
    def add(logged, site, typ, code, detail, start=None, mins=None, sup=""):
        nonlocal eid
        eid += 1
        rows.append((eid, logged, site, typ, code, detail, start or "", mins if mins is not None else "", sup))
        return eid
    ids = {}
    for (logged, site, typ, code, detail) in BENCH:
        add(logged, site, typ, code, detail)
    for r in placed:
        d, t0 = r["date"], r["start"]
        # the alarm handler stamps the field from the local clock; the operator interface stamps it
        # from the UTC clock and the export appends the house offset either way
        sd, ss = build.norm(d, r["stamp_start"])
        start_local = build.loc(sd, ss)
        logged = build.utc(d, t0 + 90)
        i = add(logged, r["site"], "reserve_unavailable", r["code"], r["reason"], start_local, r["minutes"])
        ids[r["idx"]] = i
    for r in placed:
        if not r["superseded"]:
            continue
        d = r["date"]
        add(build.utc(d, min(sim.DAY - 1, r["start"] + 7200)), r["site"], "correction", "COR-CANC",
            "event logged in error, reserve was not withdrawn", "", "", ids[r["idx"]])
    rows.sort(key=lambda x: (x[1], x[0]))
    return docs.write_csv(path, ["event_id", "logged_utc", "site_id", "event_type", "event_code",
                                 "detail", "start_local", "duration_min", "supersedes_event_id"], rows)


def register(path):
    wb = Workbook()
    ws = wb.active; ws.title = "Sites"
    hdr = ["Site id", "Site name", "Technology", "Connection (kV)", "Installed (MW)",
           "Contracted FCR (MW)", "Metering point (EAN)", "Usual FCR blocks (indicative)",
           "Controller", "Commissioned", "Location"]
    ws.append(hdr)
    for s in sim.SITES:
        ean = build.METER_EAN[s["id"]]
        ws.append([s["id"], s["name"], s["kind"].replace("_", " "), VOLTS[s["id"]], s["pmax"],
                   s["cap"], ean[1] or ean[0], ", ".join(str(b + 1) for b in s["blocks"]),
                   CONTROLLER[s["kind"]], COMMISSIONED[s["id"]], PLACE[s["id"]]])
    for c in ws[1]:
        c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFGHIJK", (11, 15, 16, 13, 13, 16, 22, 24, 26, 13, 26)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ns = wb.create_sheet("Notes")
    for line in [["Asset register extract, reserve portfolio desk"],
                 ["Pulled 2026-10-12 from the asset master. One row per reserve-providing unit."],
                 ["Metering point shows the code in force at the pull date; replaced codes are kept in the"],
                 ["metering history held by the settlement team."],
                 ["Usual FCR blocks is indicative only. The blocks a site is committed to on a delivery date"],
                 ["are in the commitment file the platform writes each day."],
                 ["Installed capacity is the connection capacity, not the capacity offered to the reserve"],
                 ["market."]]:
        ns.append(line)
    ns.column_dimensions["A"].width = 100
    wb.properties.creator = "Veerkracht Flex"; wb.properties.lastModifiedBy = "Veerkracht Flex"
    wb.properties.created = dt.datetime(2026, 10, 12, 8, 41, 0); wb.properties.modified = dt.datetime(2026, 10, 12, 8, 44, 0)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    docs.normalise_zip(path)


def metered_totals(W):
    """Net MWh per site over the window, from the same arrays the metered file is written from."""
    out = {}
    for s in sim.SITES:
        tot = 0.0
        for d in sim.dates():
            tot += float(core.metered(W, d, s["id"]).sum())
        out[s["id"]] = tot
    return out


def settlement_files(tgt, W):
    """Daily settlement statements per metering point, in two delivery months, with the October
    statement issued preliminary and then amended."""
    daily = {}
    for s in sim.SITES:
        for d in sim.dates():
            v = float(core.metered(W, d, s["id"]).sum())
            if s["id"] in HV_SITES:
                v *= LOSS_ADJ
            ean = build.METER_EAN[s["id"]]
            code = ean[0]
            if ean[1] and d >= ean[2]:
                code = ean[1]
            daily[(code, d)] = round(v, 3)
    sep = [(c, d.isoformat(), "%.3f" % v) for (c, d), v in sorted(daily.items()) if d.month == 9]
    octo = [(c, d.isoformat(), "%.3f" % v) for (c, d), v in sorted(daily.items()) if d.month == 10]
    gap_code = build.METER_EAN[PRELIM_GAP[0]][0]
    prelim = [r for r in octo if not (r[0] == gap_code and r[1] == PRELIM_GAP[1].isoformat())]
    hdr = ["metering_point", "delivery_date", "settled_volume_mwh"]
    n1 = docs.write_csv(os.path.join(tgt, "operator", "settlement_statement_daily_2026-09.csv"), hdr, sep)
    n2 = docs.write_csv(os.path.join(tgt, "operator", "settlement_statement_daily_2026-10_preliminary.csv"), hdr, prelim)
    n3 = docs.write_csv(os.path.join(tgt, "operator", "settlement_statement_daily_2026-10_amended.csv"), hdr, octo)
    stmt = {}
    for s in sim.SITES:
        ean = build.METER_EAN[s["id"]]
        tot = sum(v for (c, d), v in daily.items() if c in (ean[0], ean[1]))
        stmt[s["id"]] = tot
    lost = sum(v for (c, d), v in daily.items() if c == build.METER_EAN["VKF-OOS"][0])
    gapv = daily[(gap_code, PRELIM_GAP[1])]
    return dict(rows=(n1, n2, n3), statement=stmt, ean_lost_mwh=lost, prelim_gap_mwh=gapv)


def _block_of(sec):
    return int(sec // 14400) + 1


def availability(W, placed, read="golden"):
    """Reserve capacity hours committed and held available, per delivery block.

    read "golden"  : every live window at the place it actually ran.
    read "cracker" : clause 4 executed on the start_local field as it stands, so the rows the
                     operator interface wrote are charged to the block two hours earlier, and
                     only where the site was committed to that block on that date.
    read "mirror"  : the cracker's reading with the correction rows not applied.
    """
    committed = {b: 0.0 for b in range(1, 7)}
    for d in sim.dates():
        for s in sim.SITES:
            for b in sim.commitments(s, d):
                committed[b + 1] += s["cap"] * 4.0
    lost = {b: 0.0 for b in range(1, 7)}
    for r in placed:
        if r["superseded"] and read != "mirror":
            continue
        s = sim.SITE_BY_ID[r["site"]]
        mwh = s["cap"] * r["minutes"] / 60.0
        if read == "golden":
            lost[_block_of(r["start"])] += mwh
        else:
            d, ss = build.norm(r["date"], r["stamp_start"])
            if not (sim.WINDOW_START <= d <= sim.WINDOW_END):
                continue
            b = _block_of(ss)
            if (b - 1) in sim.commitments(s, d):
                lost[b] += mwh
    return dict(committed=committed, lost=lost,
                available={b: committed[b] - lost[b] for b in range(1, 7)})


def adherence(W, read="golden"):
    """Energy delivered above and below programme, per site, over the settlement periods of the
    window the site held no delivery-block commitment for.

    read "golden"  : each metered interval paired with the programme period it actually covers,
                     which is the period before the one its row is indexed with.
    read "natural" : each row paired with the programme period its own settlement_period names.
    """
    dates = sim.dates()
    met = {}
    for d in dates:
        for s in sim.SITES:
            met[(d, s["id"])] = core.metered(W, d, s["id"])
    above = {s["id"]: 0.0 for s in sim.SITES}
    below = {s["id"]: 0.0 for s in sim.SITES}
    for i, d in enumerate(dates):
        for s in sim.SITES:
            sid = s["id"]
            prog = W[d]["prog"][sid]
            comm = np.zeros(96, dtype=bool)
            for b in sim.commitments(s, d):
                lo, hi = sim.BLOCK_HOURS[b]
                comm[lo * 4:hi * 4] = True
            for p in range(96):
                if comm[p]:
                    continue
                if read == "golden":
                    e = met[(d, sid)][p]
                else:
                    # the row indexed with this period carries the interval that closed when it
                    # opened, which is the period before it
                    if p == 0:
                        if i == 0:
                            continue            # no row is published for it
                        e = met[(dates[i - 1], sid)][95]
                    else:
                        e = met[(d, sid)][p - 1]
                dv = float(e) - float(prog[p]) * 0.25
                if dv >= 0:
                    above[sid] += dv
                else:
                    below[sid] += -dv
    return dict(above=above, below=below,
                above_total=sum(above.values()), below_total=sum(below.values()))


def distractors(tgt):
    r = sim._rng("imbalance")
    rows = []
    for d in sim.dates():
        for p in range(96):
            h = p / 4.0
            base = 62 + 38 * np.sin(np.pi * (h - 5) / 14.0) ** 2 - 22 * np.exp(-((h - 13) ** 2) / 6.0)
            up = base + abs(r.normal(26, 17)); dn = base - abs(r.normal(24, 15))
            rows.append((build.loc(d, p * 900), p + 1, round(float(max(dn, -180)), 2),
                         round(float(min(up, 520)), 2), round(float(base + r.normal(0, 9)), 2)))
    n1 = docs.write_csv(os.path.join(tgt, "market", "imbalance_prices_15min.csv"),
                        ["period_start_local", "settlement_period", "downward_price_eur_mwh",
                         "upward_price_eur_mwh", "mid_price_eur_mwh"], rows)
    rows = []
    for d in sim.dates():
        for b in range(6):
            lo, hi = sim.BLOCK_HOURS[b]
            cleared = 11.5 + 7.5 * (b in (1, 4)) + float(sim._rng("auc", d.isoformat(), b).normal(0, 2.1))
            vol = sum(s["cap"] for s in sim.SITES if b in sim.commitments(s, d))
            rows.append((d.isoformat(), b + 1, "%02d:00" % lo, "%02d:00" % hi,
                         round(max(cleared, 3.0), 2), round(vol, 2), 112.0))
    n2 = docs.write_csv(os.path.join(tgt, "market", "fcr_capacity_auction_results.csv"),
                        ["delivery_date", "block", "block_start_local", "block_end_local",
                         "clearing_price_eur_mw_h", "awarded_mw_group", "procured_mw_total"], rows)
    return n1, n2


def prequal(tgt):
    rows = []
    for s in sim.SITES:
        r = sim._rng("prq", s["id"])
        d0 = dt.date.fromisoformat(COMMISSIONED[s["id"]]) + dt.timedelta(days=int(20 + r.randint(0, 40)))
        rows.append((s["id"], d0.isoformat(), "step test, 200 mHz, both directions",
                     round(float(s["cap"] / sim.SPAN), 2), round(float(s["cap"] / sim.SPAN * (0.98 + r.random_sample() * 0.05)), 2),
                     "pass", "witnessed by Rijksnet"))
        if s["id"] in ("VKF-KPV", "VKF-ZWT", "VKF-TRB", "VKF-HGZ"):
            d1 = dt.date(2026, 9, 16) if s["id"] in ("VKF-KPV", "VKF-ZWT") else dt.date(2026, 9, 18)
            rows.append((s["id"], d1.isoformat(), "controller change, declaration only",
                         round(float(s["cap"] / sim.SPAN), 2), "", "declared",
                         "provider declaration under clause 3, no witnessed test"))
    return docs.write_csv(os.path.join(tgt, "prequalification", "prequalification_records.csv"),
                          ["site_id", "test_date", "test_method", "contracted_gain_mw_per_hz",
                           "measured_gain_mw_per_hz", "outcome", "note"], rows)


def dictionary(tgt, counts):
    rows = [
        ("platform/vkf_highres_archive.sqlite", "recording", "one high resolution recording the platform kept",
         "recording_id, site_id, trigger, start_utc, end_utc, samples"),
        ("platform/vkf_highres_archive.sqlite", "power_1s", "net active power at the connection point, one row per second",
         "recording_id, ts_utc, p_mw, quality (M measured, E estimated)"),
        ("platform/site_programme_15min.csv", "", "the programme setpoint the platform held for each settlement period",
         "site_id, period_start_local, settlement_period, programme_mw"),
        ("platform/fcr_block_commitments.csv", "", "the delivery blocks each site was committed to, by delivery date",
         "delivery_date, block, block_start_local, block_end_local, site_id, committed_mw"),
        ("platform/control_platform_event_log.csv", "", "events the control platform logged, including reserve withdrawals and their corrections",
         "event_id, logged_utc, site_id, event_type, event_code, detail, start_local, duration_min, supersedes_event_id"),
        ("operator/system_frequency_1s_extracts.csv", "", "published system frequency at one second, one extract per deviation",
         "extract_id, ts_utc, frequency_hz"),
        ("operator/system_frequency_minute_summary.csv", "", "published minute summary of system frequency for the window",
         "minute_utc, frequency_min_hz, frequency_mean_hz, frequency_max_hz"),
        ("settlement/metered_volumes_15min.csv", "", "one interval of one metering point as the desk's meter data platform registered it, net at the connection point",
         "metering_point, interval_ts_local, settlement_period, net_volume, unit"),
        ("operator/settlement_statement_daily_2026-09.csv", "", "the operator's settled volumes per metering point and delivery date",
         "metering_point, delivery_date, settled_volume_mwh"),
        ("operator/settlement_statement_daily_2026-10_preliminary.csv", "", "as above, preliminary issue for October",
         "metering_point, delivery_date, settled_volume_mwh"),
        ("operator/settlement_statement_daily_2026-10_amended.csv", "", "as above, amended issue for October",
         "metering_point, delivery_date, settled_volume_mwh"),
        ("register/asset_register.xlsx", "Sites", "one row per reserve-providing unit in the portfolio",
         "site id, technology, connection, installed, contracted FCR, metering point, usual blocks, controller"),
        ("prequalification/prequalification_records.csv", "", "prequalification tests and controller-change declarations",
         "site_id, test_date, test_method, contracted_gain_mw_per_hz, measured_gain_mw_per_hz, outcome, note"),
        ("market/imbalance_prices_15min.csv", "", "imbalance settlement prices for the window",
         "period_start_local, settlement_period, downward_price_eur_mwh, upward_price_eur_mwh, mid_price_eur_mwh"),
        ("market/fcr_capacity_auction_results.csv", "", "daily FCR capacity auction outcomes and the group's awarded volume",
         "delivery_date, block, clearing_price_eur_mw_h, awarded_mw_group, procured_mw_total"),
    ]
    return docs.write_csv(os.path.join(tgt, "data_dictionary.csv"),
                          ["file", "table_or_sheet", "one row is", "columns"], rows)
