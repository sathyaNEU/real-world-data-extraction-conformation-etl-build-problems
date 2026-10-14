"""task457's assertions, run on every build: the input gates; the golden from the bundle against the generator's truth;
every rung by name; the whole estimator family with its margins; every semantic-fork row; the three-rules block on the
clean truth; mess neutrality per Layer 1 item; the asks' goldens, stops and cracker; the separation of the asks from
the main call; the distractors' deletion test; the bins; the realism sweep."""
import collections
import csv
import datetime as dt
import json
import os
import random
import re

import asks as A
import conform as CF
import golden as GD
import methods as M
import pack as P
import world as W

RESULTS = []


def ok(label, cond, detail=""):
    RESULTS.append((label, bool(cond), detail))


def pdf_text(path):
    from pypdf import PdfReader
    return "\n".join(p.extract_text() or "" for p in PdfReader(path).pages)


class TruthCF(M.Conformed):
    """The conformed data straight from the generator's world: every relationship-year as paid (with its payments' dates)
    and as first disclosed, the day each disclosure was first submitted, with no record attached to the wrong person,
    no amendment, no typed name and no vintage."""

    def __init__(self, fac):
        import people as PE
        self.pub, self.dis, self.pay, self.day = {}, {}, {}, {}
        self.fac = {f.emp_id: None for f in fac}
        self.filed = set()
        self._mcache, self._ocache, self._fw = {}, {}, {}
        for f in fac:
            for y, t in getattr(f, "filing", {}).items():
                if t is not None:
                    self.filed.add((f.emp_id, y))
                    self.day[(f.emp_id, y)] = PE.eastern_day(t)
            for rel in f.rels:
                if rel["mfr"] is None:
                    continue
                for y, a in rel["years"].items():
                    if y <= 2025:
                        self.pub[(f.emp_id, rel["mfr"], y)] = self.pub.get((f.emp_id, rel["mfr"], y), 0.0) + a
                        self.pay.setdefault((f.emp_id, rel["mfr"], y), []).extend(rel["pay"][y])
                    d = f.lines.get((rel["mfr"], y))
                    if d and (f.emp_id, y) in self.filed:
                        self.dis[(f.emp_id, rel["mfr"], y)] = d
        self.by_person_year = collections.defaultdict(dict)
        for (e, m, y), d in self.dis.items():
            self.by_person_year[(e, y)][m] = d
        self.pub_by_person_year = collections.defaultdict(dict)
        for (e, m, y), a in self.pub.items():
            self.pub_by_person_year[(e, y)][m] = a


NATURAL = ([("line", "face", 0.0), ("line", "pool", 0.0), ("all", "own", 0.0), ("line", "own", 0.0)] +
           [("all", "usd", K) for K in M.GRIDS["usd"]] + [("all", "n", K) for K in M.GRIDS["n"]] +
           [("line", "n", K) for K in M.GRIDS["n"]] + [("line", "usd", K) for K in M.GRIDS["usd"]])


def default_path(cf):
    """The analyst's route: the natural corrections (face value, the faculty-wide ratio, each faculty member's overall
    share or matched-line ratio, own or weighted toward the faculty's), every disclosure read as covering its whole year,
    compared on 2025 by the filed criterion, the closest applied."""
    scored = sorted((M.backtest(cf, g, h, K), g, h, K) for g, h, K in NATURAL)
    best = scored[0]
    dk, ranked = M.docket(cf, best[1], best[2], best[3])
    return best, dk, ranked, scored


def run(target, golden_dir, world=None, meta_path=None):
    RESULTS.clear()
    data = CF.load(target)
    cf = M.Conformed(data)
    fac_w = world["fac"]
    grp = {f.emp_id: f.group for f in fac_w}
    members = lambda kinds: {f.emp_id for f in fac_w if f.group in kinds}
    G = M.golden(cf)
    gold = set(G["docket"])
    label = lambda dk: " ".join("%s%d" % kv for kv in sorted(collections.Counter(grp[e] for e in dk).items()))

    # ------------------------------------------------------------------ input gates
    files = sorted(os.listdir(target))
    meta = json.load(open(meta_path)) if meta_path else {}
    ok("gate: 10 or more files", len(files) >= 10, len(files))
    ok("gate: 3 or more formats", len({f.rsplit(".", 1)[1] for f in files}) >= 3, sorted({f.rsplit(".", 1)[1] for f in files}))
    ok("gate: one file of 25,000 rows or more", len(data["register"]) >= 25000, len(data["register"]))
    ok("gate: two or more distractors named in metadata.json and shipped",
       len(meta.get("distractor_files", [])) >= 2 and all(d in files for d in meta.get("distractor_files", [])))
    ok("gate: no distractor named as such in a file name", not any("distractor" in f for f in files))

    # ------------------------------------------------------------------ the golden
    ok("golden: the docket is the five large faithful disclosers, the five long-record under-reporters and the five "
       "December filers whose retainers' December installments came after their disclosures",
       gold == members({"F", "U", "C"}), label(gold))
    ok("golden: the weighting chosen on 2025 is the grid's closest (K=%s matched lines)" % G["K"],
       G["K"] == min(M.GRIDS["n"], key=lambda K: M.backtest(cf, "line", "n", K, "sum", align="both")))
    gap = G["ranked"][14][1][0] / G["ranked"][15][1][0]
    ok("golden: the fifteenth stands at least 1.25 times the sixteenth", gap >= 1.25, round(gap, 3))
    truth = TruthCF(fac_w)
    filed26 = {d["emp_id"] for d in data["disc"] if int(d["disclosure_year"]) == 2026}
    paid26 = {}
    for f in fac_w:
        if f.emp_id in filed26 and f.owes(2026):
            v = [rel["years"][2026] for rel in f.rels if rel["mfr"] and 2026 in rel["years"]
                 and f.lines.get((rel["mfr"], 2026))]
            if v:
                paid26[f.emp_id] = max(v)
    top_truth = set(sorted(paid26, key=lambda e: -paid26[e])[:15])
    ok("golden: the docket is the fifteen largest 2026 single-manufacturer remunerations the manufacturers paid "
       "(the generator's truth, unpublished)", top_truth == gold, label(top_truth))
    an = world["golden"]
    ok("golden: golden.py's docket is the checks' docket", {m["emp_id"] for m in an["members"][:15]} == gold)
    dec_gold = {e for e in gold if cf.day[(e, 2026)].year == 2026}
    ok("golden: the docket's December filers are the five C members", dec_gold == members({"C"}), label(dec_gold))

    # ------------------------------------------------------------------ rungs
    face = set(M.docket(cf, "line", "face")[0])
    ok("R0 face value: the under-reporters and the December filers partly give way", face != gold, label(face))
    ok("R0' the faculty-wide ratio: another docket", set(M.docket(cf, "line", "pool")[0]) != gold)
    best, dk, ranked, scored = default_path(cf)
    ok("R3 the default path's comparison on 2025 picks a matched-line ratio weighted toward the faculty's, every disclosure "
       "read as its whole year, and it files the five censored-history faculty in place of the five December filers",
       best[1] == "line" and best[2] in ("n", "usd") and set(dk) == members({"F", "U", "X"}),
       "%s %s" % (best[1:], label(dk)))
    ok("R3 the default path's choice beats face value by 25 per cent or more on 2025",
       best[0] <= 0.75 * M.backtest(cf, "line", "face"), round(best[0] / M.backtest(cf, "line", "face"), 3))
    all_own = set(M.docket(cf, "all", "own")[0])
    ok("R1 each faculty member's overall share: the omitters join", members({"O"}) <= all_own and all_own != gold,
       label(all_own))
    for h, K in (("usd", 2e4), ("n", 3), ("yrs", 2)):
        s = set(M.docket(cf, "all", h, K)[0])
        ok("R1' the overall share weighted by %s (K=%s): the omitters join" % (h, K), len(members({"O"}) & s) >= 3 and s != gold,
           label(s))
    line_own = set(M.docket(cf, "line", "own")[0])
    ok("R2 the matched-line ratio unweighted: the thin records join", members({"T"}) <= line_own and line_own != gold,
       label(line_own))
    hist_only = set(M.docket(cf, "line", "n", G["K"], align="hist")[0])
    target_only = set(M.docket(cf, "line", "n", G["K"], align="target")[0])
    ok("R4 either half of the alignment alone (the earlier years counted to their submission, or the 2026 December "
       "disclosures grossed up) files the golden docket at the golden weighting", hist_only == gold and target_only == gold,
       "%s / %s" % (label(hist_only), label(target_only)))
    face_al = set(M.docket(cf, "line", "face", align="both")[0])
    ok("R4' the December disclosures grossed up with no reporting correction: the under-reporters stay off",
       face_al != gold and not members({"U"}) <= face_al, label(face_al))

    # ------------------------------------------------------------------ the family
    fam = M.family(cf, aligns=M.ALIGNS, covers=M.COVERS)
    hyb = M.hybrid_rows(cf, aligns=M.ALIGNS, covers=M.COVERS)
    gold_t = tuple(sorted(gold))
    weighted = lambda r: (r[1] == "line" or (r[1] == "pair" and r[5] == "line")) and r[2] not in ("face", "pool", "own")
    unal = [r for r in fam if r[8] == "none"]
    un_best = {}
    for r in unal:
        k = (r[1], r[2], r[4], r[5])
        un_best[k] = min(un_best.get(k, r[0]), r[0])
    un_chosen = [r for r in unal if r[0] == un_best[(r[1], r[2], r[4], r[5])]]
    ok("family: read as covering its whole year, every construction at its chosen weighting files another docket "
       "(%d constructions)" % len(un_chosen), all(r[6] != gold_t for r in un_chosen),
       [r[1:6] for r in un_chosen if r[6] == gold_t][:3])
    ok("family: the closest whole-year construction files another docket and stands 1.15 times the golden or more",
       min(unal)[6] != gold_t and min(unal)[0] >= 1.15 * G["bt"], round(min(unal)[0] / G["bt"], 3))
    both_line = [r for r in fam if r[8] == "both" and weighted(r) and r[9] in ("pair", "add")]
    ok("family: every aligned matched-line construction with a short-record weighting files the golden docket (%d cells)"
       % len(both_line), all(r[6] == gold_t for r in both_line), [r[1:6] + r[8:] for r in both_line if r[6] != gold_t][:3])
    every = fam + hyb
    closer = [r for r in every if r[0] <= G["bt"]]
    ok("family+: every construction coming closer to 2025 than the golden files the golden docket (%d cells)" % len(closer),
       all(r[6] == gold_t for r in closer), [r[1:6] + r[8:] for r in closer if r[6] != gold_t][:3])
    ok("family+: the closest construction of all files the golden docket", min(every)[6] == gold_t, min(every)[1:6])
    wrong = min(r[0] for r in every if r[6] != gold_t)
    ok("family+: every construction filing another docket stands 1.08 times the golden or more (the half-aligned ones "
       "closest)", wrong >= 1.08 * G["bt"], round(wrong / G["bt"], 3))
    ok("family+: every construction filing another docket stands 1.2 times the closest construction or more",
       wrong >= 1.2 * min(every)[0], round(wrong / min(every)[0], 3))
    whole = min(r[0] for r in every if r[8] == "none")
    ok("family+: every construction reading disclosures as whole years stands 1.15 times the golden and 1.3 times the "
       "closest construction or more", whole >= 1.15 * G["bt"] and whole >= 1.3 * min(every)[0],
       "%.3f / %.3f" % (whole / G["bt"], whole / min(every)[0]))
    best_of = {}
    for r in both_line:
        k = (r[1], r[2], r[4], r[5], r[9])
        best_of[k] = min(best_of.get(k, r[0]), r[0])
    chosen = [r for r in both_line if r[0] == best_of[(r[1], r[2], r[4], r[5], r[9])]]
    near = [r for r in every if r[0] <= 1.1 * G["bt"]]
    ok("family: every construction closer to 2025 than the golden (%d cells) holds the fifteenth at least 1.13 times the "
       "sixteenth" % len(closer), min(r[7] for r in closer) >= 1.13, round(min(r[7] for r in closer), 3))
    near_full = [r for r in near if r[8] == "both"]
    ok("family: every fully aligned construction within 10 per cent of the golden on 2025 (%d cells) files the golden "
       "docket with the fifteenth at least 1.10 times the sixteenth" % len(near_full),
       all(r[6] == gold_t for r in near_full) and min(r[7] for r in near_full) >= 1.10,
       round(min(r[7] for r in near_full), 3))
    fw = [r for r in fam if r[8] == "both" and r[9] == "fw" and weighted(r)]
    fw_best = {}
    for r in fw:
        k = (r[1], r[2], r[4], r[5])
        fw_best[k] = min(fw_best.get(k, r[0]), r[0])
    fw_chosen = [r for r in fw if r[0] == fw_best[(r[1], r[2], r[4], r[5])]]
    ok("family: grossed up by the faculty-wide share instead, every construction's chosen weighting files the golden "
       "docket (%d constructions)" % len(fw_chosen), all(r[6] == gold_t for r in fw_chosen),
       [r[1:6] for r in fw_chosen if r[6] != gold_t][:3])
    umin, cmin, tmax = [], [], []
    for r in [x for x in chosen + fw_chosen if x[0] <= 1.1 * G["bt"]]:
        est = M.estimates(cf, 2026, r[1], r[2], r[3], r[4], r[5], align=r[8], cover=r[9])
        ranked_r = sorted(est.values(), key=lambda v: -v[0])
        sixteenth, fifteenth = ranked_r[15][0], ranked_r[14][0]
        umin.append(min(est[e][0] for e in members({"U"})) / sixteenth)
        cmin.append(min(est[e][0] for e in members({"C"})) / sixteenth)
        tmax.append(max(est[e][0] for e in members({"T", "X", "P", "O"})) / fifteenth)
    ok("family: within 10 per cent of the golden, every under-reporter stands at least 1.1 times the sixteenth at each "
       "chosen weighting", min(umin) >= 1.1, round(min(umin), 3))
    ok("family: within 10 per cent of the golden, every December filer stands at least 1.1 times the sixteenth at each "
       "chosen weighting and share", min(cmin) >= 1.1, round(min(cmin), 3))
    ok("family: within 10 per cent of the golden, the thin records, censored histories, faculty just below and omitters "
       "sit under the fifteenth place", max(tmax) < 1.0, round(max(tmax), 3))

    # ------------------------------------------------------------------ semantic forks
    def docket_under(cls=M.Conformed, **kw):
        c2 = cls(data, **kw)
        g2 = M.golden(c2)
        return set(g2["docket"]), c2

    d1, _ = docket_under(natures="all")
    ok("SF1 remuneration: every nature on the register side files the same docket", d1 == gold, label(d1))
    d2, _ = docket_under(personal=False)
    ok("SF2 personally: payments to the university kept files the same docket", d2 == gold)
    ok("SF3 a single manufacturer: every docket member's 2026 manufacturers are distinct Open Payments IDs",
       all(len(cf.by_person_year[(e, 2026)]) == len(set(cf.by_person_year[(e, 2026)])) for e in gold))
    ok("SF4 the calendar year: the disclosures' year is the register's program year (2026 disclosures cover calendar 2026)",
       all(int(d["disclosure_year"]) in W.YEARS for d in data["disc"]))
    region_all = {f.emp_id for f in fac_w if f.group != "bg"}
    amended_region = [d for d in data["disc"] if d["emp_id"] in region_all and d["current_version"] != 1]
    ok("SF5 the disclosure in force: no docket-region member amended a disclosure, so latest, first and summed versions agree",
       not amended_region and docket_under(version="first")[0] == gold and docket_under(version="sum")[0] == gold)
    departed = [e for e in cf.fac if (e, 2026) in cf.by_person_year and not CF.holds(cf.fac[e], dt.date(2026, 12, 31))]
    ok("SF6 faculty estimated: nobody without an appointment on 31 December 2026 carries a 2026 manufacturer line",
       not departed, departed[:3])
    d7, _ = docket_under(vintage="jun")
    ok("SF7 the latest published year: the June 2026 publication of 2025 files the same docket", d7 == gold, label(d7))
    for metric in ("rmse", "rel"):
        best_cells = []
        for r in fam:
            est = M.estimates(cf, 2025, r[1], r[2], r[3], r[4], r[5], align=r[8], cover=r[9])
            tr = M.truth(cf, 2025, est)
            if metric == "rmse":
                s = (sum((est[e][0] - tr[e]) ** 2 for e in est) / len(est)) ** 0.5
            else:
                s = sum(abs(est[e][0] - tr[e]) / max(tr[e], 1.0) for e in est) / len(est)
            best_cells.append((s, r[6]))
        ok("SF8 closest by %s: the closest construction still files the golden docket" % metric, min(best_cells)[1] == gold_t)
    est25 = M.estimates(cf, 2025, "line", "n", G["K"], "sum", align="both")
    tr_disc = {e: max(cf.pub_by_person_year.get((e, 2025), {}).get(m, 0.0) for m in cf.by_person_year[(e, 2025)]) for e in est25}
    btd = lambda al, h, K: sum(abs(M.estimates(cf, 2025, "line", h, K, "sum", align=al)[e][0] - tr_disc[e]) for e in est25)
    ok("SF9 the comparison on disclosed manufacturers only: the aligned weighting still beats every whole-year reading",
       btd("both", "n", G["K"]) < min(btd("none", h, K) for h, K in (("face", 0.0), ("n", 2), ("n", 4), ("usd", 1e4))))
    d10, _ = docket_under(attribution="matcher")
    d10b, _ = docket_under(attribution="npi")
    ok("SF10 records of the faculty: the office's matcher and NPI-only attribution file the same docket", d10 == gold and d10b == gold)
    ok("SF11 a matched line needs both sources: the docket unchanged with disclosed lines the register lacks counted at zero",
       _sf11(data, gold))
    ok("SF12 the fifteen: no tie at the fifteenth in any golden cell", min(r[7] for r in every if r[6] == gold_t) > 1.0)
    d13, _ = docket_under(cls=_UtcDay)
    ok("SF13 the submission's day: read in UTC rather than Eastern time the same docket", d13 == gold, label(d13))
    d14, _ = docket_under(cls=_StrictDay)
    ok("SF14 a payment dated the day of the submission left out: the same docket", d14 == gold, label(d14))
    d15, _ = docket_under(cls=_FirstDay)
    ok("SF15 the first submission's day instead of the version in force: the same docket", d15 == gold, label(d15))

    # ------------------------------------------------------------------ three rules on the clean truth
    tb = default_path(truth)
    ok("R3a analyst test: on the clean truth the default path's comparison picks a weighted matched-line ratio and files "
       "the censored-history faculty in place of the December filers",
       tb[0][1] == "line" and set(tb[1]) == members({"F", "U", "X"}), "%s %s" % (tb[0][1:], label(tb[1])))
    tg = M.golden(truth)
    ok("R3b the decisive step alone: the clean truth with each disclosure counted to its submission files the golden docket",
       set(tg["docket"]) == gold, label(tg["docket"]))
    ok("R3b the grain alone on the clean truth (aligned, matched lines unweighted) stays wrong",
       set(M.docket(truth, "line", "own", align="both")[0]) != gold)
    ok("R3b the weighting alone on the clean truth (aligned overall share weighted) stays wrong",
       all(set(M.docket(truth, "all", h, K, align="both")[0]) != gold for h, K in (("usd", 2e4), ("n", 3), ("yrs", 2))))
    ok("R3b the alignment alone on the clean truth (no reporting correction) stays wrong",
       set(M.docket(truth, "line", "face", align="both")[0]) != gold)
    cols = {f: hdr for f, hdr in _headers(target).items()}
    banned = re.compile(r"(ratio|complete|estimate|credib|weight|share|omit|understat|flag|is_|coverage|received)", re.I)
    ok("R1 the raw material ships: payment records with their dates and disclosure versions with their submission times",
       len(data["register"]) > 50000 and len(data["items"]) > 3000 and len(data["history"]) > 5000 and G["docket"])
    ok("R1 no shipped column holds a ratio, an estimate, a weight, a share received or a flag",
       not [(f, c) for f, hs in cols.items() for c in hs if banned.search(c)],
       [(f, c) for f, hs in cols.items() for c in hs if banned.search(c)][:4])
    docs_text = pdf_text(os.path.join(target, P.FILES["policy"])) + pdf_text(os.path.join(target, P.FILES["guide"]))
    names_g = [m["last"] for m in an["members"][:15]]
    ok("R3c no plant: no document names a docket member, a ratio, a weighting, an estimator or a disclosure's reach",
       not any(n in docs_text for n in names_g) and not re.search(
           r"\b(credib\w*|shrink\w*|weighted|ratio|matched|installment\w*|retainer\w*|so far|to date|received by|"
           r"payments? (?:made|received) (?:after|before))\b", docs_text, re.I))
    for item, kw in (("the office's matcher", {"attribution": "matcher"}), ("records carrying an NPI only", {"attribution": "npi"}),
                     ("payments to the university kept", {"personal": False}), ("every nature", {"natures": "all"}),
                     ("the first version", {"version": "first"}), ("every version summed", {"version": "sum"}),
                     ("typed names left out", {"typed": "list"}), ("typed names against present names only", {"typed": "present"}),
                     ("the June 2026 publication", {"vintage": "jun"})):
        d, _ = docket_under(**kw)
        ok("R2 %s: the golden method files the same docket" % item, d == gold, label(d))
    rows = list(data["register"])
    random.Random(457).shuffle(rows)
    shuffled = {k: v for k, v in data.items() if not isinstance(k, tuple) and k not in ("faculty",)}
    shuffled["register"] = rows
    ok("R2 row order: the register shuffled files the same docket", set(M.golden(M.Conformed(shuffled))["docket"]) == gold)

    # ------------------------------------------------------------------ asks
    st = A.stops(data)
    a, b, c = st["A"]["golden"], st["B"]["golden"], st["C"]["golden"]
    ok("A golden equals the docket's table", a == an["asks"]["A"]["golden"])
    need_a = {"employee numbers joined as they stand": 16, "latest submission time": 16, "UTC date against 15 January": 12,
              "roster's department code": 4, "everyone on the roster at the pull": 8}
    for name, v in st["A"].items():
        if name == "golden":
            continue
        moved = sum(1 for k in set(a) | set(v) if a.get(k) != v.get(k))
        ok("A stop '%s' moves %d of %d departments" % (name, moved, len(a)), moved >= need_a[name])
    for name, v in st["B"].items():
        if name == "golden":
            continue
        moved = sum(1 for k in b if GD.thousand(b[k]) != GD.thousand(v[k]))
        ok("B stop '%s' moves %d of 5 natures" % (name, moved), moved >= (5 if name in ("the office's matched employee number", "every record in the extract") else 1))
    need_c = {"the list's blank identifier trusted": 4, "typed names against the list's present names": 3,
              "every typed name": 3, "the first version": 1, "every version": 1}
    for name, v in st["C"].items():
        if name == "golden":
            continue
        moved = sum(1 for k in c if GD.thousand(c[k]) != GD.thousand(v[k]))
        ok("C stop '%s' moves %d of 4 types" % (name, moved), moved >= need_c[name])
    for k, v in b.items():
        edge = abs((v % 1000) - 500)
        ok("bin: B %s sits %d dollars from a rounding edge" % (A.NATURE_LABEL[k][:20], edge), edge >= 70, round(v, 2))
    for k, v in c.items():
        edge = abs((v % 1000) - 500)
        ok("bin: C %s sits %d dollars from a rounding edge" % (k, edge), edge >= 70, round(v, 2))
    ok("bin: B total sits 70 dollars or more from a rounding edge", abs((sum(b.values()) % 1000) - 500) >= 70)
    ok("bin: C total sits 70 dollars or more from a rounding edge", abs((sum(c.values()) % 1000) - 500) >= 70)

    # ------------------------------------------------------------------ separation of the asks from the main call
    region = region_all
    att_m = CF.attribute(data, "matcher")
    att_p = CF.attribute(data, "profile")
    wrong_att = [rid for rid, e in att_m.items() if att_p.get(rid) != e and e in region]
    ok("separation: no other physician's record is attached to a docket-region member", not wrong_att, len(wrong_att))
    blank_region = [r for r in data["register"] if att_p.get(r["Record_ID"]) in region and not r["Covered_Recipient_NPI"]]
    ok("separation: no docket-region record lacks its NPI", not blank_region, len(blank_region))
    third_region = [r for r in data["register"] if att_p.get(r["Record_ID"]) in region and r["Third_Party_Payment_Recipient_Indicator"] != "No Third Party Payment"]
    ok("separation: no docket-region payment was made to a third party", not third_region)
    dids = {d["disclosure_id"]: d["emp_id"] for d in data["disc"]}
    typed_region = [it for it in data["items"] if dids[it["disclosure_id"]] in region and not it["entity_id"]]
    ok("separation: every docket-region item names a listed company", not typed_region, len(typed_region))
    late_region = [e for e in members({"F", "P", "U", "O", "T", "C", "X"}) if CF.dept_on(data, cf.fac[e], dt.date(2026, 12, 31)) is None]
    ok("separation: every docket-region member holds a primary appointment on 31 December 2026 under a current code",
       not late_region and all(t[2] != "MED" for e in members({"F", "P", "U", "O", "T", "C", "X"}) for t in cf.fac[e]["terms"]
                               if t[1] >= dt.date(2026, 12, 31)))
    rekeyed = [f for f in fac_w if getattr(f, "rekey_day", None)]
    ok("separation: the faculty whose employee number changed hold no industry relationship and none is in the docket region",
       rekeyed and all(not f.rels and f.group == "bg" for f in rekeyed), len(rekeyed))
    nowrep_ids = {m["id"] for m in world["mfrs"] if m.get("now_reporting")}
    nr_people = {e for (e, m, y) in cf.dis if m in nowrep_ids}
    ok("separation: no line with a company that began reporting is a faculty member's largest, and none is in the region",
       all(cf.by_person_year[(e, 2026)][m] < max(v for mm, v in cf.by_person_year[(e, 2026)].items() if mm not in nowrep_ids)
           for (e, m, y) in cf.dis if m in nowrep_ids) and not (nr_people & region))
    cf_list = M.Conformed(data, listed="list")
    ok("separation: read through the list's blank identifiers, the estimated faculty and the docket are unchanged",
       {e for (e, y) in cf_list.by_person_year if y == 2026} == {e for (e, y) in cf.by_person_year if y == 2026}
       and set(M.golden(cf_list)["docket"]) == gold)

    # ------------------------------------------------------------------ cracker: the readings a careful analyst takes
    crack_a = A.ask_a(data, link="emp")
    ok("cracker A: employee numbers joined as they stand misses every department's late count",
       sum(1 for k in a if a[k][1] != crack_a[k][1]) == len(a))
    crack_b = A.ask_b(data, attribution="matcher")
    ok("cracker B: the office's matched column trusted misses every nature", all(GD.thousand(b[k]) != GD.thousand(crack_b[k]) for k in b))
    crack_c = A.ask_c(data, listed="list")
    ok("cracker C: the list's blank identifiers trusted miss every type", all(GD.thousand(c[k]) != GD.thousand(crack_c[k]) for k in c))

    # ------------------------------------------------------------------ distractors
    used = {"register", "profiles", "roster", "depts", "entities", "export", "history"}
    ok("distractors: the golden reads none of the research extract, the NIH awards or the June 2026 publication",
       not ({"research", "nih", "register_june"} & used) and all(P.FILES[k] in meta.get("distractor_files", [])
                                                                    for k in ("research", "nih", "register_june")))

    # ------------------------------------------------------------------ realism sweep
    amts = [float(it["amount_usd"]) for it in data["items"] if it["amount_usd"] not in (None, "")]
    top_amt = max(collections.Counter(amts).values())
    ok("realism: self-reported amounts heap on round figures as reports do (a third or more on whole $500s, no one amount "
       "above 15 per cent)", sum(1 for x in amts if x % 500 == 0) >= len(amts) / 3 and top_amt <= 0.15 * len(amts), top_amt)
    amounts = [float(r["Total_Amount_of_Payment_USDollars"]) for r in data["register"]]
    ok("realism: register amounts are not round (under 6 per cent whole hundreds)",
       sum(1 for x in amounts if x % 100 == 0) < 0.06 * len(amounts), round(sum(1 for x in amounts if x % 100 == 0) / len(amounts), 3))
    ok("realism: ask A's department counts all differ", len({v[0] for v in a.values()}) >= 14)
    ok("realism: the personas are the card's", all(p in (W.PEOPLE.values()) for p in ("Karen Shea", "Diana Hickman", "Herbert Dean")))
    rid = [int(r["Record_ID"]) for r in data["register"]]
    diffs = collections.Counter(y - x for x, y in zip(sorted(rid), sorted(rid)[1:]))
    ok("realism: Record_IDs advance irregularly", max(diffs.values()) < 0.03 * len(rid))
    dates = [dt.datetime.strptime(r["Date_of_Payment"], "%m/%d/%Y").date() for r in data["register"]]
    ok("realism: no register date after the extract", max(dates) <= dt.date(2025, 12, 31))
    hist_t = [dt.datetime.strptime(e["event_time_utc"], "%Y-%m-%dT%H:%M:%SZ") for e in data["history"]]
    ok("realism: no submission after the export's pull", max(hist_t) <= dt.datetime(2027, 1, 21, 15, 0))
    secs = collections.Counter(t.second for t in hist_t)
    ok("realism: submission times carry no favoured second", max(secs.values()) < 0.04 * len(hist_t), max(secs.values()))
    emp_bg = sorted(int(f.emp_id) for f in fac_w if f.group == "bg")
    ok("realism: the docket region's employee numbers sit inside the roster's own range",
       all(emp_bg[0] < int(e) < emp_bg[-1] for e in region))
    d26 = sorted(cf.by_person_year[(e, 2026)][m] for e in gold for m in [G["ranked"][[x[0] for x in G["ranked"]].index(e)][1][1]])
    steps = collections.Counter(round(y - x, -2) for x, y in zip(d26, d26[1:]))
    ok("realism: the docket's disclosed amounts follow no arithmetic progression", max(steps.values()) <= 3, steps.most_common(2))
    addr = collections.Counter(r["Recipient_Primary_Business_Street_Address_Line1"] for r in data["register"])
    ok("realism: no one business address on more than 40 per cent of the register's rows",
       max(addr.values()) < 0.4 * len(data["register"]), addr.most_common(1))
    prods = {r["Name_of_Drug_or_Biological_or_Device_or_Medical_Supply_1"] for r in data["register"]} - {""}
    ok("realism: product names are words, not codes", not [p for p in prods if re.search(r"\d", p)])
    figs = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "figures.json")))
    ok("figures.json carries the docket", sorted(figs["call"]) == sorted(m["name"] for m in an["members"][:15]))

    failed = [r for r in RESULTS if not r[1]]
    for lbl, good, det in RESULTS:
        print("%s  %s%s" % ("PASS" if good else "FAIL", lbl, ("  [%s]" % det) if det not in ("", None) and not good else ""))
    print("checks: %d of %d pass" % (len(RESULTS) - len(failed), len(RESULTS)))
    return not failed


class _ZeroPublished(M.Conformed):
    """A disclosed manufacturer-year the register lacks taken as a line with nothing published."""

    def matched(self, e, before, aligned=False):
        key = (e, before, aligned)
        if key in self._mcache:
            return self._mcache[key]
        out = []
        for y in range(2021, before):
            for m, d in self.by_person_year.get((e, y), {}).items():
                a = self.pub_by_person_year.get((e, y), {}).get(m, 0.0)
                if aligned and a:
                    a = self.to_day(e, m, y)
                if d:
                    out.append((a, d, m, y))
        self._mcache[key] = out
        return out


class _UtcDay(M.Conformed):
    """The submission's day read in UTC."""

    def __init__(self, data, **kw):
        super().__init__(data, **kw)
        self.day = dict(self.day)
        for d in data["disc"]:
            v = d["last_submitted_utc"]
            t = v if isinstance(v, dt.datetime) else dt.datetime.strptime(v, "%Y-%m-%d %H:%M:%S")
            self.day[(d["emp_id"], int(d["disclosure_year"]))] = t.date()


class _StrictDay(M.Conformed):
    """A payment dated the day of the submission not yet received."""

    def __init__(self, data, **kw):
        super().__init__(data, **kw)
        self.day = {k: v - dt.timedelta(days=1) for k, v in self.day.items()}


class _FirstDay(M.Conformed):
    """The day the disclosure was first submitted, whatever version is in force."""

    def __init__(self, data, **kw):
        super().__init__(data, **kw)
        self.day = dict(CF.in_force_day(data, "first"))


def _sf11(data, gold):
    cf = _ZeroPublished(data)
    return set(M.golden(cf)["docket"]) == gold


def _headers(target):
    from openpyxl import load_workbook
    out = {}
    for f in sorted(os.listdir(target)):
        p = os.path.join(target, f)
        if f.endswith(".csv"):
            with open(p, encoding="utf-8") as fh:
                out[f] = next(csv.reader(fh))
        elif f.endswith(".xlsx"):
            wb = load_workbook(p, read_only=True)
            out[f] = sorted({str(x) for ws in wb.worksheets for x in next(ws.iter_rows(max_row=1, values_only=True))})
    return out
