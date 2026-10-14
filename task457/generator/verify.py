"""An independent check of task457's goldens from the bundle alone, on its own code path (pandas): the docket under
the matched-line weighting chosen on 2025 with each disclosure counted to the day it was submitted, the whole-year
reading's docket, and every figure of the three supporting tables, each compared with the deliverables."""
import datetime as dt
import os
import sys

import numpy as np
import pandas as pd

REMUN = {"Consulting Fee",
         "Compensation for services other than consulting, including serving as faculty or as a speaker at a venue other "
         "than a continuing education program",
         "Honoraria",
         "Compensation for serving as faculty or as a speaker for an accredited or certified continuing education program",
         "Compensation for serving as faculty or as a speaker for an unaccredited and non-certified continuing education "
         "program"}
TYPES = {"Consulting", "Speaking", "Advisory board", "Honoraria", "Other paid services", "Consulting or advisory board",
         "Speaking or teaching", "Expert testimony"}


def load(t):
    reg = pd.read_csv(os.path.join(t, "open_payments_general_2021-2025_jan2027_extract.csv"), dtype=str, keep_default_na=False)
    prof = pd.read_csv(os.path.join(t, "open_payments_covered_recipient_profiles.csv"), dtype=str, keep_default_na=False)
    ros = pd.read_csv(os.path.join(t, "clinical_faculty_appointments.csv"), dtype=str, keep_default_na=False)
    ent = pd.read_csv(os.path.join(t, "coi_entity_list.csv"), dtype=str, keep_default_na=False)
    disc = pd.read_excel(os.path.join(t, "coi_disclosures_2021-2026_export.xlsx"), sheet_name="Disclosures", dtype=str)
    items = pd.read_excel(os.path.join(t, "coi_disclosures_2021-2026_export.xlsx"), sheet_name="Items", dtype=str).fillna("")
    hist = pd.read_csv(os.path.join(t, "coi_submission_history.csv"), dtype=str)
    depts = pd.read_csv(os.path.join(t, "department_codes.csv"), dtype=str, keep_default_na=False)
    return reg, prof, ros, ent, disc, items, hist, depts


def conformed(reg, prof, ros, ent, disc, items):
    npi_emp = ros.drop_duplicates("npi").set_index("npi")["emp_id"]
    pnpi = prof.set_index("Covered_Recipient_Profile_ID")["Covered_Recipient_NPI"]
    r = reg[reg["Nature_of_Payment_or_Transfer_of_Value"].isin(REMUN)
            & (reg["Third_Party_Payment_Recipient_Indicator"] == "No Third Party Payment")].copy()
    r["emp"] = r["Covered_Recipient_Profile_ID"].map(pnpi).map(npi_emp)
    r = r.dropna(subset=["emp"])
    r["amt"] = r["Total_Amount_of_Payment_USDollars"].astype(float)
    r["op"] = r["Applicable_Manufacturer_or_Applicable_GPO_Making_Payment_ID"]
    r["year"] = r["Program_Year"].astype(int)
    r["date"] = pd.to_datetime(r["Date_of_Payment"], format="%m/%d/%Y")
    pay = r[["emp", "op", "year", "date", "amt"]]
    pub = r.groupby(["emp", "op", "year"])["amt"].sum().reset_index()
    names = dict(zip(reg["Applicable_Manufacturer_or_Applicable_GPO_Making_Payment_Name"].str.strip().str.lower(),
                     reg["Applicable_Manufacturer_or_Applicable_GPO_Making_Payment_ID"]))
    names.update(dict(zip(ent.loc[ent["open_payments_id"] != "", "entity_name"].str.strip().str.lower(),
                          ent.loc[ent["open_payments_id"] != "", "open_payments_id"])))
    eid = dict(zip(ent["entity_id"], ent["open_payments_id"]))
    it = items.merge(disc[["disclosure_id", "emp_id", "disclosure_year", "current_version"]], on="disclosure_id")
    it = it[(it["version"] == it["current_version"]) & it["relationship_type"].isin(TYPES)].copy()
    listed = it["entity_id"].map(eid).fillna("")
    by_name = it["entity_name"].str.strip().str.lower().map(names).fillna("")
    it["op"] = np.where(listed != "", listed, by_name)
    it = it[it["op"] != ""]
    it["amt"] = it["amount_usd"].astype(float)
    dis = it.groupby(["emp_id", "op", "disclosure_year"])["amt"].sum().reset_index()
    dis.columns = ["emp", "op", "year", "dis"]
    dis["year"] = dis["year"].astype(int)
    m = dis.merge(pub.rename(columns={"amt": "pub"}), on=["emp", "op", "year"], how="left")
    d = disc.copy()
    d["day"] = (pd.to_datetime(d["last_submitted_utc"]) - pd.Timedelta(hours=5)).dt.normalize()
    d["year"] = d["disclosure_year"].astype(int)
    day = d.set_index(["emp_id", "year"])["day"]
    m["day"] = [day.get((e, y)) for e, y in zip(m["emp"], m["year"])]
    m["dec"] = [x.year == y for x, y in zip(m["day"], m["year"])]
    got = pay.merge(m[["emp", "op", "year", "day"]], on=["emp", "op", "year"])
    got = got[got["date"] <= got["day"]].groupby(["emp", "op", "year"])["amt"].sum()
    m["pub_to_day"] = [(got.get((e, o, y), 0.0) if dec else p) for e, o, y, dec, p in
                       zip(m["emp"], m["op"], m["year"], m["dec"], m["pub"])]
    return m, pub, pay


def estimate(m, pub, pay, year, K, aligned=True):
    col = "pub_to_day" if aligned else "pub"
    hist = m[(m["year"] < year) & (m["pub"] > 0) & (m["dis"] > 0) & (m[col] > 0)]
    pool = hist["dis"].sum() / hist[col].sum()
    g = hist.groupby("emp").agg(d=("dis", "sum"), p=(col, "sum"), n=("dis", "size"))
    g["c"] = g["n"] / (g["n"] + K) * g["d"] / g["p"] + K / (g["n"] + K) * pool
    cur = m[m["year"] == year].copy()
    cur["c"] = cur["emp"].map(g["c"]).fillna(pool)
    cur["q"] = 1.0
    if aligned:
        prev = pay[pay["year"] == year - 1].copy()
        prev["md"] = prev["date"].dt.month * 100 + prev["date"].dt.day
        allprev = prev["amt"].sum()
        for i in cur.index[cur["dec"]]:
            md = cur.at[i, "day"].month * 100 + cur.at[i, "day"].day
            pp = prev[(prev["emp"] == cur.at[i, "emp"]) & (prev["op"] == cur.at[i, "op"])]
            tot, hit = pp["amt"].sum(), pp.loc[pp["md"] <= md, "amt"].sum()
            cur.at[i, "q"] = hit / tot if tot > 0 and hit > 0 else prev.loc[prev["md"] <= md, "amt"].sum() / allprev
    cur["est"] = cur["dis"] / cur["c"] / cur["q"]
    return cur.sort_values(["emp", "est"], ascending=[True, False]).groupby("emp").head(1).set_index("emp")


def docket(m, pub, pay, K, aligned=True):
    e = estimate(m, pub, pay, 2026, K, aligned)
    top = e.sort_values("est", ascending=False)
    return list(top.index[:15]), top


def backtest(m, pub, pay, K, aligned=True):
    e = estimate(m, pub, pay, 2025, K, aligned)
    tr = pub[pub["year"] == 2025].groupby("emp")["amt"].max()
    return float((e["est"] - tr.reindex(e.index).fillna(0.0)).abs().sum())


def main(target, golden_dir):
    reg, prof, ros, ent, disc, items, hist, depts = load(target)
    m, pub, pay = conformed(reg, prof, ros, ent, disc, items)
    grid = [0.5, 1, 1.5, 2, 3, 4, 6, 8, 12, 20]
    scores = {K: backtest(m, pub, pay, K) for K in grid}
    K = min(scores, key=scores.get)
    plain = {K2: backtest(m, pub, pay, K2, aligned=False) for K2 in grid}
    K_plain = min(plain, key=plain.get)
    dk, top = docket(m, pub, pay, K)
    pdk, _ = docket(m, pub, pay, K_plain, aligned=False)
    names = ros.drop_duplicates("emp_id").set_index("emp_id")
    golden_names = sorted("%s %s" % (names.loc[e, "first_name"], names.loc[e, "last_name"]) for e in dk)
    from docx import Document
    doc = Document(os.path.join(golden_dir, "coi_committee_docket_2027-02.docx"))
    table_names = sorted(row.cells[0].text.replace(" *", "") for row in doc.tables[0].rows[1:])
    checks = []
    import json
    figs = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "figures.json")))
    checks.append(("verify: the weighting chosen on 2025 is the one the deliverable used", K == figs["weighting_k_lines"], K))
    checks.append(("verify: counted to each submission, the chosen weighting comes closer to 2025 than every whole-year "
                   "weighting", scores[K] < min(plain.values()), (round(scores[K]), round(min(plain.values())))))
    checks.append(("verify: the docket's fifteen match the deliverable", golden_names == table_names, golden_names[:3]))
    checks.append(("verify: the whole-year reading files a different fifteen", set(pdk) != set(dk), len(set(pdk) & set(dk))))
    checks.append(("verify: the fifteenth stands clear of the sixteenth", top["est"].iloc[14] / top["est"].iloc[15] > 1.13,
                   round(top["est"].iloc[14] / top["est"].iloc[15], 3)))
    # ask A: a faculty member is a person, the employee numbers sharing an NPI one faculty member
    ros["s"] = pd.to_datetime(ros["term_start"])
    ros["e"] = pd.to_datetime(ros["term_end"])
    day = pd.Timestamp("2026-12-31")
    on = ros[(ros["s"] <= day) & (ros["e"] >= day)]
    prim = on[on["appointment_type"] == "Primary"].drop_duplicates("npi").set_index("npi")
    succ = dict(zip(depts["formed_from"], depts["dept_code"]))
    dept = prim.apply(lambda r: succ["MED/" + r["division_code"]] if r["dept_code"] == "MED" else r["dept_code"], axis=1)
    owed = set(on["npi"])
    emp_npi = ros.drop_duplicates("emp_id").set_index("emp_id")["npi"]
    d26 = disc[disc["disclosure_year"] == "2026"].set_index("disclosure_id")["emp_id"]
    h = hist[(hist["event"] == "Submitted") & hist["disclosure_id"].isin(d26.index)].copy()
    h["t"] = pd.to_datetime(h["event_time_utc"].str.replace("Z", ""))
    h["npi"] = h["disclosure_id"].map(d26).map(emp_npi)
    first = h.groupby("npi")["t"].min()
    deadline = pd.Timestamp("2027-01-16 04:59:59")
    late = {p for p in owed if p not in first.index or first[p] > deadline}
    tA = {row.cells[0].text: (int(row.cells[1].text.replace(",", "")), int(row.cells[2].text.replace(",", "")))
          for row in doc.tables[1].rows[1:-1]}
    code_name = {r["dept_code"]: r["dept_name"] for _, r in depts.iterrows()}
    vA = {}
    for p in sorted(owed):
        nm = code_name[dept[p]]
        o, l = vA.get(nm, (0, 0))
        vA[nm] = (o + 1, l + (p in late))
    checks.append(("verify: ask A's sixteen departments match the deliverable", vA == tA, [(k, vA.get(k), tA.get(k)) for k in tA if vA.get(k) != tA.get(k)][:3]))
    # ask B
    npi_emp = ros.drop_duplicates("npi").set_index("npi")["emp_id"]
    pnpi = prof.set_index("Covered_Recipient_Profile_ID")["Covered_Recipient_NPI"]
    r = reg[(reg["Program_Year"] == "2025") & reg["Nature_of_Payment_or_Transfer_of_Value"].isin(REMUN)
            & (reg["Third_Party_Payment_Recipient_Indicator"] == "No Third Party Payment")].copy()
    r["emp"] = r["Covered_Recipient_Profile_ID"].map(pnpi).map(npi_emp)
    held = set(ros[(ros["s"] <= pd.Timestamp("2025-12-31")) & (ros["e"] >= pd.Timestamp("2025-01-01"))]["emp_id"])
    r = r[r["emp"].isin(held)]
    vB = r.assign(a=r["Total_Amount_of_Payment_USDollars"].astype(float)).groupby("Nature_of_Payment_or_Transfer_of_Value")["a"].sum()
    tB = [int(row.cells[1].text.replace("$", "").replace(",", "")) for row in doc.tables[2].rows[1:-1]]
    order = ["Consulting Fee", "Compensation for services other than consulting, including serving as faculty or as a speaker "
             "at a venue other than a continuing education program", "Honoraria",
             "Compensation for serving as faculty or as a speaker for an accredited or certified continuing education program",
             "Compensation for serving as faculty or as a speaker for an unaccredited and non-certified continuing education program"]
    checks.append(("verify: ask B's five natures match the deliverable", [int(round(vB[k] / 1000.0)) * 1000 for k in order] == tB, tB))
    # ask C
    it = items.merge(disc[["disclosure_id", "disclosure_year", "current_version"]], on="disclosure_id")
    it = it[(it["disclosure_year"] == "2026") & (it["version"] == it["current_version"]) & it["relationship_type"].isin(TYPES)].copy()
    names_op = set(reg["Applicable_Manufacturer_or_Applicable_GPO_Making_Payment_Name"].str.strip().str.lower()) | set(
        ent.loc[ent["open_payments_id"] != "", "entity_name"].str.strip().str.lower())
    eid = dict(zip(ent["entity_id"], ent["open_payments_id"]))
    listed_op = it["entity_id"].map(eid).fillna("") != ""
    nr = it[~listed_op & ~it["entity_name"].str.strip().str.lower().isin(names_op)]
    vC = nr.assign(a=nr["amount_usd"].astype(float)).groupby("relationship_type")["a"].sum()
    tC = {row.cells[0].text: int(row.cells[1].text.replace("$", "").replace(",", "")) for row in doc.tables[3].rows[1:-1]}
    checks.append(("verify: ask C's four types match the deliverable",
                   all(int(round(vC.get(k, 0.0) / 1000.0)) * 1000 == v for k, v in tC.items()), tC))
    bad = [c for c in checks if not c[1]]
    for c in checks:
        print("%s  %s%s" % ("PASS" if c[1] else "FAIL", c[0], ("  [%s]" % (c[2],)) if not c[1] and len(c) > 2 else ""))
    print("verify: %d of %d pass" % (len(checks) - len(bad), len(checks)))
    if bad:
        raise SystemExit("verify failed")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
