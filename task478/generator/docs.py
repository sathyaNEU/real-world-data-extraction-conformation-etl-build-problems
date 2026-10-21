"""The pack's documents and workbooks: the Spring Patching Program 2027 (PDF), Kerbline's pilot scope of work (PDF),
the City's route turnback notice (PDF), the data notes (PDF), the review module export and the unit register (xlsx)."""
from reportlab import rl_config
rl_config.invariant = 1
from reportlab.lib import colors                                     # noqa: E402
from reportlab.lib.pagesizes import letter                           # noqa: E402
from reportlab.lib.styles import ParagraphStyle                      # noqa: E402
from reportlab.lib.units import inch                                 # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402
from openpyxl import Workbook                                        # noqa: E402
from openpyxl.styles import Font                                     # noqa: E402

import pack as P    # noqa: E402
import world as W   # noqa: E402

BODY = ParagraphStyle("b", fontName="Times-Roman", fontSize=10.5, leading=13.6, spaceAfter=6)
SMALL = ParagraphStyle("sm", fontName="Times-Roman", fontSize=9.5, leading=12, spaceAfter=4)
HEAD = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=11, leading=14, spaceBefore=8, spaceAfter=4)
TITLE = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=15, leading=19, spaceAfter=4)
SUB = ParagraphStyle("s", fontName="Helvetica", fontSize=9, leading=12, textColor=colors.HexColor("#444444"),
                     spaceAfter=10)
MONO = ParagraphStyle("m", fontName="Courier", fontSize=8.6, leading=10.6, spaceAfter=1)


def _doc(path, title, org):
    d = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.95 * inch, rightMargin=0.95 * inch,
                          topMargin=0.9 * inch, bottomMargin=0.9 * inch, title=title, author=org,
                          creator=org, producer=org, subject=title)

    def frame(c, d_):
        c.saveState()
        c.setFont("Helvetica", 8)
        c.setFillColor(colors.HexColor("#555555"))
        c.drawString(0.95 * inch, 0.55 * inch, "%s  |  %s" % (org, title))
        c.drawRightString(letter[0] - 0.95 * inch, 0.55 * inch, "Page %d" % d_.page)
        c.restoreState()
    return d, frame


ORG = "%s, %s" % (W.DOT, W.DISTRICT)


def write_program(path):
    title = "Spring Patching Program 2027"
    d, frame = _doc(path, title, ORG)
    eng = W.PEOPLE["engineer"]
    s = [Paragraph("Spring Patching Program 2027", TITLE),
         Paragraph("%s. Maintenance. Issued 6 April 2027 by %s, P.E., District Maintenance Engineer." % (ORG, eng), SUB),
         Paragraph("1. Purpose", HEAD),
         Paragraph("Each spring the district puts hot-mix patching crews to work under its on-call pavement repair "
                   "contract (contract 26-OC-1187) to make permanent repairs where the winter has done the most damage "
                   "to the state roadway. For 2027 the district gives a dedicated crew, for May and June, to each "
                   "maintenance section whose roadway came out of March in the worst condition. The other sections are "
                   "covered by their own section crews and the contractor's shared crew.", BODY),
         Paragraph("2. Which sections get a crew", HEAD),
         Paragraph("2.1 A maintenance section receives a dedicated crew when its state-maintained roadway held 6.0 or "
                   "more potholes per ten lane-miles in March 2027.", BODY),
         Paragraph("2.2 The count is every pothole on that roadway in March 2027, whether or not a Kerbline unit "
                   "recorded it. A pothole that was cold patched in March counts: the spring crews replace cold patch "
                   "with hot mix.", BODY),
         Paragraph("2.3 A pothole is a bowl-shaped hole in the traveled way at least six inches across (Pavement "
                   "Distress Manual, section 4.3). The district reviews every defect Kerbline reports as a pothole and "
                   "rejects any that is not one.", BODY),
         Paragraph("2.4 A section's roadway is its state-maintained segments in the district road inventory as at 31 "
                   "March 2027. Lane-miles are each segment's length times its through lanes.", BODY),
         Paragraph("3. Task orders", HEAD),
         Paragraph("The task orders issue on 19 April 2027 and the crews start on 3 May 2027. A designated section's "
                   "task order names its yard as the crew's base and lists the section's state routes. Each task order "
                   "also lists the section's March cold patches, for replacement, and the miles the pilot's camera units "
                   "were driven on the section's roadway in March.", BODY),
         Paragraph("4. The March 2027 detection campaign", HEAD),
         Paragraph("From 1 to 31 March 2027 the district ran Kerbline Vision camera units on its section patrol pickups, "
                   "freeway service patrol trucks, street sweepers and district crew trucks under the pilot scope of "
                   "work. Pothole reports went to the district's review queue, and the section crews cold patched from "
                   "the queue through the month.", BODY),
         Spacer(1, 10),
         Paragraph("%s, P.E.<br/>District Maintenance Engineer, %s" % (eng, W.DISTRICT), BODY)]
    d.build(s, onFirstPage=frame, onLaterPages=frame)


def write_sow(path):
    title = "Pothole Detection Pilot: Scope of Work"
    org = "Kerbline Vision, Inc."
    d, frame = _doc(path, title, org)
    s = [Paragraph("Tamsett District Pothole Detection Pilot: Scope of Work", TITLE),
         Paragraph("Kerbline Vision, Inc. Attachment A to purchase order 27-PO-0418, %s. January 2027." % ORG, SUB),
         Paragraph("1. Services", HEAD),
         Paragraph("Kerbline supplies camera units, installs them on up to 36 district vehicles the district chooses, and "
                   "provides its platform, its review module and data exports for the pilot, which runs from 22 February "
                   "to 30 June 2027.", BODY),
         Paragraph("2. Units", HEAD),
         Paragraph("A unit is a smartphone in a windshield mount running the Kerbline detection app. It powers with its "
                   "vehicle and runs a session on the platform while the vehicle's ignition is on, recording the roadway "
                   "ahead and reporting the defects it detects: potholes, alligator cracking, longitudinal cracking, "
                   "utility covers, patches, edge breaks and debris.", BODY),
         Paragraph("3. Defect numbers", HEAD),
         Paragraph("Kerbline links repeated detections of one defect into one defect number. The detection inventory "
                   "lists the defects the units detected; it is not an inventory of every defect on the roadway the "
                   "vehicles drove.", BODY),
         Paragraph("4. Review", HEAD),
         Paragraph("Every pothole-class defect goes to the district's review queue, where the district's reviewer "
                   "confirms or rejects it. Kerbline supports the reviewer's work in the review module.", BODY),
         Paragraph("5. Data", HEAD),
         Paragraph("On request Kerbline provides the detections export, with each detection's capture time in district "
                   "local time; the review module export; and the unit session export from its device API, with times "
                   "in UTC.", BODY),
         Paragraph("6. Charges", HEAD),
         Paragraph("6.1 A unit's service day is a calendar day, district local time, during any part of which the unit, "
                   "installed on a district vehicle, had a session open on the platform.", BODY),
         Paragraph("6.2 The district pays $%.2f per service day, invoiced monthly in arrears. Bench testing before "
                   "installation is not charged." % W.DAY_RATE, BODY),
         Paragraph("6.3 Installation, removal, the review module and data exports are included.", BODY),
         Paragraph("7. Term", HEAD),
         Paragraph("Either party may end the pilot on 30 days' notice. Units are returned to Kerbline at the end of the "
                   "pilot or when the district takes a vehicle out of the pilot.", BODY)]
    d.build(s, onFirstPage=frame, onLaterPages=frame)


def write_turnback(path, segs):
    title = "Route Turnback: Notice of Effective Date"
    org = W.CITY
    d, frame = _doc(path, title, org)
    rows = [["Route", "Segment", "From MP", "To MP", "Length (mi)", "Through lanes"]]
    for s in sorted([s for s in segs if s["turnback"]], key=lambda s: (s["route"], s["bmp"])):
        rows.append([s["route"], str(s["seg"]), "%.3f" % s["bmp"], "%.3f" % s["emp"], "%.3f" % s["length"], str(s["lanes"])])
    t = Table(rows, hAlign="LEFT", colWidths=[0.8 * inch, 0.9 * inch, 0.8 * inch, 0.8 * inch, 0.9 * inch, 1.0 * inch])
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 9), ("FONT", (0, 1), (-1, -1), "Helvetica", 9),
                           ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.black), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    s = [Paragraph("State Route Turnback: Notice of Effective Date", TITLE),
         Paragraph("%s, Department of Public Works, and %s. 1 March 2027." % (W.CITY, ORG), SUB),
         Paragraph("Under the turnback agreement of 12 November 2026 between the %s and the %s, jurisdiction over the "
                   "state highway segments listed below passes to the City at 12:01 a.m. on 15 March 2027. From that time "
                   "the City maintains them, including pavement repair, sweeping and winter operations, and the district's "
                   "maintenance sections no longer patrol them as state roadway." % (W.DOT, W.CITY), BODY),
         Spacer(1, 6), t, Spacer(1, 10),
         Paragraph("Segment numbers are the district road inventory's 2026 release. Questions go to the City Engineer's "
                   "office or the district's planning engineer.", BODY)]
    d.build(s, onFirstPage=frame, onLaterPages=frame)


def write_dictionary(path, n):
    title = "Pilot data notes, March 2027 extracts"
    d, frame = _doc(path, title, ORG)
    F = P.FILES
    items = [
        (F["detections"], n["detections"],
         "One row per detection: a unit's capture of a defect. The platform keeps the first capture of a defect by a "
         "unit on a calendar day. defect_id is the platform's defect number. road_ref is the road inventory segment the "
         "platform's base map gives the detection (the base map is the 2025 inventory release); road_mp is the route "
         "milepost; both are blank off the state highway system. captured_local is district local time. severity (1 to 3) "
         "is Kerbline's for pothole-class defects."),
        (F["traversals"], n["traversals"],
         "One row per traversal of a state highway segment by a pilot vehicle, from the fleet telematics provider's map "
         "matching. segment_id is the 2026 inventory release. enter_utc and exit_utc are UTC."),
        (F["review"], n["review"],
         "One row per pothole-class defect with the reviewer's decision (Confirmed or Rejected) and, for a rejection, the "
         "reason as entered. Milepost is the defect's route milepost. Times are UTC."),
        (F["register"], n["register"],
         "One row per unit: the vehicle it is installed on, the fleet group, and the installation and removal dates."),
        (F["sessions"], n["sessions"],
         "Per unit, its platform sessions from Kerbline's device API: session start and end (UTC), frames and megabytes "
         "uploaded."),
        (F["sections"], n["sections"],
         "The district's maintenance sections: number, name and yard."),
        (F["inventory"], n["inventory"],
         "The 2026 release (as at 1 January 2027): one row per state highway segment in the district's fourteen "
         "maintenance sections, with its route mileposts, length and through lanes. A segment keeps its id from release "
         "to release unless it is split or merged."),
        (F["crosswalk"], n["crosswalk"],
         "The 2025 segments that the 2026 release split or merged, with the 2026 segments they became."),
        (F["orders"], n["orders"],
         "One row per task line of a cold patch work order (activity PCOLD) raised in March. ext_ref is the Kerbline defect "
         "number the order was raised from. raised and finished are district local time. cold_mix_lb is the cold mix "
         "pounds as entered. status is COMPLETE or OPEN at the extract."),
        (F["pci"], n["pci"],
         "The 2026 pavement condition survey: one row per segment, with the survey date, roughness (IRI), rut depth, "
         "cracking and the pavement condition index."),
        (F["winter"], n["winter"],
         "The March 2027 storm events by section: plow hours, salt, brine and freeze-thaw cycles."),
    ]
    s = [Paragraph("Tamsett District pothole detection pilot: data notes", TITLE),
         Paragraph("%s. Maintenance Management Systems, %s. Compiled 3 April 2027." % (ORG, W.PEOPLE["analyst"]), SUB)]
    for name, rows, text in items:
        s.append(Paragraph("<b>%s</b> (%s rows)" % (name, "{:,}".format(rows)), BODY))
        s.append(Paragraph(text, SMALL))
    s.append(Paragraph("Times: the district's local time is Eastern Time; daylight saving time began at 2:00 a.m. on 14 March "
                       "2027.", BODY))
    d.build(s, onFirstPage=frame, onLaterPages=frame)


def write_review(path, e):
    wb = Workbook()
    ws = wb.active
    ws.title = "Review decisions"
    ws.append(["Defect ID", "Class", "Route", "Milepost", "First detected (UTC)", "Decision", "Reason", "Reviewed (UTC)",
               "Reviewer"])
    for c in ws[1]:
        c.font = Font(bold=True)
    n = 0
    for row in P.review_rows(e):
        ws.append(row)
        n += 1
    for r_ in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        for c in r_:
            c.number_format = "yyyy-mm-dd hh:mm"
    for r_ in ws.iter_rows(min_row=2, min_col=8, max_col=8):
        for c in r_:
            c.number_format = "yyyy-mm-dd hh:mm"
    for r_ in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        for c in r_:
            c.number_format = "0.000"
    for col, w in zip("ABCDEFGHI", (12, 9, 22, 10, 19, 11, 30, 17, 9)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:I%d" % (n + 1)
    wb.save(path)
    return n


def write_register(path, e):
    wb = Workbook()
    ws = wb.active
    ws.title = "Units"
    ws.append(["Unit serial", "Vehicle", "Fleet group", "Vehicle description", "Home section", "Installed", "Removed",
               "Mount", "Notes"])
    for c in ws[1]:
        c.font = Font(bold=True)
    n = 0
    for row in P.register_rows(e):
        ws.append(row)
        n += 1
    for col in ("F", "G"):
        for c in ws[col][1:]:
            c.number_format = "yyyy-mm-dd"
    for col, w in zip("ABCDEFGHI", (11, 10, 22, 36, 13, 12, 12, 19, 48)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(path)
    return n
