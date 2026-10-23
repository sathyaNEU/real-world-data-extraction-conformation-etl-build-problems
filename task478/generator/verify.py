"""An independent recompute of task478's graded figures from the cut pack, sharing no code with bundle.py: each
facility's 2026 NOx from its Rule 4306 units, the five, the gas Tamsett Valley Gas delivered and the emergency engines'
NOx; then the same figures read back from figures.json, submission.md and the three goldens.

    python3 task478/generator/verify.py
"""
import csv
import datetime as dt
import json
import os
import re
import sys
import warnings

import numpy as np
import pandas as pd

warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
T = os.path.join(TASK, "target")
G = os.path.join(TASK, "golden")
OK = []


def say(name, ok):
    OK.append(bool(ok))
    print("%s  %s" % ("PASS" if ok else "FAIL", name))


def half_up(x):
    return int(np.floor(x + 0.5))


def main():
    p = lambda fn: os.path.join(T, fn)
    pm = pd.read_csv(p("tbapcd_permit_devices.csv"), dtype=str)
    fac_name = pm.drop_duplicates("FAC_NO").set_index("FAC_NO")["FACILITY"].to_dict()
    order = list(pd.read_csv(os.path.join(G, "tbapcd_nox_inventory_2026.csv"), dtype=str)["facility"])
    u = pm[pm["RULE"] == "4306"].set_index("DEVICE_NO")
    fac_of = u["FACILITY"].to_dict()

    # the heat content table and each facility's BTU area, read off its own bills' factors
    hc = pd.read_csv(p("tvg_monthly_heat_content_2018_2027.csv"), dtype=str)
    H = {(r.BTU_AREA, r.MONTH): float(r.BTU_FACTOR) for r in hc.itertuples()}
    b = pd.read_csv(p("tvg_commercial_billing_extract_2027-02.csv"), dtype=str)
    for c in ("READ_FROM", "READ_TO"):
        b[c] = pd.to_datetime(b[c], format="%m/%d/%Y")
    b["usage"] = b["USAGE_THERMS"].astype(float)
    b["btu"] = b["BTU_FACTOR"].astype(float)
    import pypdf
    plan = re.sub(r"\s+", " ", " ".join(pg.extract_text() for pg in pypdf.PdfReader(
        p("tbapcd_lowry_creek_cerp_2027.pdf")).pages))
    norm = lambda a: re.sub(r"[^A-Z0-9]", "", re.sub(
        r"\b(ROAD|AVENUE|STREET|BOULEVARD|LANE|DRIVE|WAY)\b", lambda m: {"ROAD": "RD", "AVENUE": "AVE", "STREET": "ST",
                                                                         "BOULEVARD": "BLVD", "LANE": "LN", "DRIVE": "DR",
                                                                         "WAY": "WY"}[m.group(1)], a.upper()))
    addr = {}
    for f in order:
        m = re.search(re.escape(f) + r" \d{4} (\d+ [A-Z][\w. ]+?(?:Road|Avenue|Street|Boulevard|Lane|Drive|Way|Row))",
                      plan)
        addr[norm(m.group(1))] = f
    b["fac"] = [addr[norm(a)] for a in b["SERVICE_ADDRESS"]]
    area = {}
    for f, g in b.groupby("fac"):
        hits = {a for a in ("WESTSIDE", "EASTSIDE")
                if all(abs(H[(a, "%d-%02d" % (t.year, t.month))] - x) < 1e-9 for t, x in zip(g["READ_TO"], g["btu"]))}
        assert len(hits) == 1, f
        area[f] = hits.pop()

    # monitored units: the hourly mass, joined through the monitoring plan's unit id
    cems = pd.read_csv(p("tbapcd_cems_hourly_2026.csv"), dtype={"facility_id": str, "unit_id": str,
                                                                "nox_rate_modc": str})
    plan_id = u[u["MONITORING"] == "CEMS"].reset_index()[["FAC_NO", "CEMS_UNIT_ID", "DEVICE_NO"]]
    cems = cems.merge(plan_id, left_on=["facility_id", "unit_id"], right_on=["FAC_NO", "CEMS_UNIT_ID"], how="left")
    assert cems["DEVICE_NO"].notna().all()
    lb = cems.groupby("DEVICE_NO")["nox_mass_lb"].sum().to_dict()

    # the monitored low-NOx burner's rise against its area's heat content, from its measured hours
    burner = [d for d in plan_id["DEVICE_NO"] if "burner" in u.loc[d, "NOX_CONTROL"]]
    assert len(burner) == 1
    pb = cems[(cems["DEVICE_NO"] == burner[0]) & (cems["nox_rate_modc"] == "01")]
    xa = np.array([H[(area[fac_of[burner[0]]], x[:7])] / 1.035 - 1 for x in pb["op_date"]])
    rise = float(np.polyfit(xa, np.log(pb["nox_rate_lb_mmbtu"].astype(float)), 1)[0])

    # unmonitored units: tests in lb/MMBtu, 2026 monthly heat input
    o = pd.read_csv(p("tbapcd_source_tests_2018_2021.csv"), dtype=str)
    tests = pd.DataFrame({"dev": "D-" + o["DEVICE_NO"], "date": pd.to_datetime(o["TEST_DATE"], format="%m/%d/%Y"),
                          "lb": o["NOX_PPMVD"].astype(float) * 1.194e-7 * 8710.0 * 20.9 /
                          (20.9 - o["O2_REF"].astype(float))})
    n = pd.read_excel(p("tbapcd_source_tests_2022_2026.xlsx"), sheet_name="Results")
    tests = pd.concat([tests, pd.DataFrame({"dev": n["Device No."], "date": pd.to_datetime(n["Test Date"]),
                                            "lb": n["NOx (lb/MMBtu)"].astype(float)})], ignore_index=True)
    st = pd.read_csv(p("tbapcd_emission_statements_2018_2026.csv"), dtype=str)
    st = st[st["STMT_YEAR"] == "2026"].copy()
    q = st["QUANTITY"].astype(float)
    sc = [H[(area[fac_of[d]], "2026-%02d" % int(m))] * 1000.0 for d, m in zip(st["DEVICE_NO"], st["MONTH"])]
    st["mmbtu"] = np.select([st["UNIT"] == "THERMS", st["UNIT"] == "MMBTU", st["UNIT"] == "MMSCF"],
                            [q / 10, q, q * np.array(sc)])
    open_units = [d for d in u.index if u.loc[d, "MONITORING"] != "CEMS"]
    for d in open_units:
        a = area[fac_of[d]]
        g = tests[tests["dev"] == d]
        lv = np.exp(np.mean(np.log(g["lb"]) - rise * np.array([H[(a, t.strftime("%Y-%m"))] / 1.035 - 1
                                                               for t in g["date"]])))
        s = st[st["DEVICE_NO"] == d]
        lb[d] = float(sum(r.mmbtu * lv * np.exp(rise * (H[(a, "2026-%02d" % int(r.MONTH))] / 1.035 - 1))
                          for r in s.itertuples()))
    tons = {f: 0.0 for f in order}
    for d, v in lb.items():
        tons[fac_of[d]] += v / 2000.0
    five = sorted(sorted(order, key=lambda f: -tons[f])[:5], key=order.index)

    # ask A: corrected CCF before the July 2026 reads, therms from them; a combined account's usage once a period
    b["therms"] = np.where(b["READ_TO"] >= pd.Timestamp(2026, 7, 1), b["usage"], b["usage"] * b["btu"])
    same = b.groupby(["ACCOUNT_NO", "READ_TO"]).agg(n=("METER_NO", "nunique"), k=("usage", "nunique"))
    comb = {a for a, g in same.reset_index().groupby("ACCOUNT_NO") if (g["n"] > 1).all() and (g["k"] == 1).all()}
    b = b[~(b["ACCOUNT_NO"].isin(comb) & b.duplicated(["ACCOUNT_NO", "READ_TO"]))]
    y0, y1 = pd.Timestamp(2026, 1, 1), pd.Timestamp(2027, 1, 1)
    share = ((np.minimum(b["READ_TO"], y1) - np.maximum(b["READ_FROM"], y0)).dt.days.clip(lower=0) /
             (b["READ_TO"] - b["READ_FROM"]).dt.days)
    gas = (b["therms"] * share / 10.0).groupby(b["fac"]).sum().to_dict()

    # ask B: every log as hours and minutes, bank lines at the outage, replaced engines
    xl = pd.ExcelFile(p("tbapcd_rule4702_engine_logs_2026.xlsx"))
    logs = []
    for sh in xl.sheet_names:
        df = pd.read_excel(xl, sheet_name=sh, header=1)
        df.columns = ["date", "dev", "v", "why"]
        df["fac"] = fac_name[sh]
        logs.append(df)
    logs = pd.concat(logs, ignore_index=True)
    logs["date"] = pd.to_datetime(logs["date"])
    frac = lambda v: np.round((v - np.floor(v + 1e-9)) * 100)
    minutes = bool((frac(logs["v"]) <= 59).all())
    hm = lambda v: np.floor(v + 1e-9) + frac(v) / 60.0
    ev = pd.read_csv(p("tbapcd_compliance_events_2018_2026.csv"), dtype=str)
    ev["d"] = [pd.Timestamp(dt.datetime.strptime(x, "%m/%d/%Y")) if "/" in x else pd.Timestamp(x)
               for x in ev["EVENT_DATE"]]
    outage = {(fac_name[r.FAC_NO], r.d): float(re.search(r"for ([\d.]+) hours", r.DESCRIPTION).group(1))
              for r in ev[ev["EVENT_TYPE"] == "OUTAGE"].itertuples()}
    old = {}
    for r in ev[ev["EVENT_TYPE"] == "EQUIPMENT"].itertuples():
        m = re.search(r"removed .*?\(([\d,]+) bhp, (Tier [^)]+)\)", r.DESCRIPTION)
        old[r.DEVICE_NO] = (r.d, float(m.group(1).replace(",", "")), m.group(2))
    tier = {"Tier 1": 6.9, "Tier 2": 4.8, "Tier 3": 3.0, "Tier 4 Final": 0.30}
    allp = pm.set_index("DEVICE_NO")
    lbs = {f: 0.0 for f in order}
    for r in logs.itertuples():
        h = hm(r.v) if minutes else r.v
        if r.why.startswith("Emergency") and (r.fac, r.date) in outage:
            h = min(h, outage[(r.fac, r.date)])
        bhp, t = float(allp.loc[r.dev, "RATING"]), allp.loc[r.dev, "ENGINE_TIER"]
        if r.dev in old and r.date < old[r.dev][0]:
            bhp, t = old[r.dev][1], old[r.dev][2]
        lbs[r.fac] += h * bhp * 0.74 * tier[t] / 453.6

    # read back
    fig = json.load(open(os.path.join(HERE, "figures.json")))
    say("the five: %s" % five, fig["call"]["facilities"] == five)
    say("the 14 whole tons", all(fig["facilities"][f]["nox_tons"] == half_up(tons[f]) for f in order))
    say("the 14 gas figures, whole MMBtu", all(fig["facilities"][f]["gas_mmbtu"] == half_up(gas[f]) for f in order))
    say("the 14 engine figures, whole pounds", all(fig["facilities"][f]["engine_nox_lb"] == half_up(lbs[f])
                                                   for f in order))
    say("the rise: %.2f in the log a unit of heat content over 1.035, about %d per cent a per cent"
        % (rise, round(100 * (np.exp(rise / 100) - 1))),
        fig["components"]["rise_pct_per_pct"] == int(round(100 * (np.exp(rise / 100) - 1))))
    say("the logs hold hours and minutes: no entry's fraction reaches .60 (%d entries)" % len(logs),
        minutes and fig["ask_b"]["minutes"] is True)
    rows = list(csv.DictReader(open(os.path.join(G, "tbapcd_nox_inventory_2026.csv"), encoding="utf-8")))
    say("the golden CSV's 14 rows", [int(r["rule_4306_nox_tons_2026"]) for r in rows] ==
        [half_up(tons[f]) for f in order])
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(G, "tbapcd_cerp_facilities_2027.xlsx"))
    ag = [r[0] for r in wb["Agreements 2027"].iter_rows(min_row=5, max_row=9, values_only=True)]
    say("the workbook's agreements sheet names the five", ag == five)
    fr = {r[0]: r for r in wb["Facilities"].iter_rows(min_row=2, values_only=True) if r[0] in order}
    say("the workbook's facilities sheet: tons, agreement, gas and engines on all 14",
        all(fr[f][2] == half_up(tons[f]) and fr[f][3] == ("Yes" if f in five else "No") and fr[f][4] == half_up(gas[f])
            and fr[f][5] == half_up(lbs[f]) for f in order))
    svg = open(os.path.join(G, "tbapcd_facility_nox_2026.svg"), encoding="utf-8").read()
    say("the chart labels every facility and its whole tons", all(f in svg for f in order) and
        all(re.search(r">%d<" % half_up(tons[f]), svg) for f in order))
    sub = open(os.path.join(TASK, "submission.md"), encoding="utf-8").read()
    say("submission.md block 4 carries every graded figure",
        all("%s: %d" % (f, half_up(tons[f])) in sub and "%s: %s" % (f, "{:,}".format(half_up(gas[f]))) in sub and
            "%s: %s" % (f, "{:,}".format(half_up(lbs[f]))) in sub for f in order))
    bad = sum(1 for x in OK if not x)
    print("%s  task478 verify: %d checks, %d failed" % ("PASS" if not bad else "FAIL", len(OK), bad))
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
