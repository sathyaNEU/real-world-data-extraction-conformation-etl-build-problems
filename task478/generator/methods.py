"""Everything the call and the asks compute, read from the cut pack: the golden, the default path and the other rungs,
the licensed variants of the golden, and the ask with its natural wrong handlings. checks.py, golden.py and ceiling.py
read the pack through this module; verify.py has its own code."""
import bisect
import csv
import datetime as dt
import json
import os
import re
from collections import defaultdict
from zoneinfo import ZoneInfo

import numpy as np

TZ = ZoneInfo("America/New_York")
UTC = dt.timezone.utc
LINE = 6.0
FT = 5280.0
UP = {"NB": 1, "EB": 1, "SB": -1, "WB": -1}      # route mileposts increase northbound and eastbound


def _csv(target, name):
    with open(os.path.join(target, name), encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def _utc(s):
    return dt.datetime.strptime(s, "%Y-%m-%dT%H:%M:%S.%fZ").replace(tzinfo=UTC)


def _local(s, fmt="%Y-%m-%d %H:%M:%S"):
    return dt.datetime.strptime(s, fmt).replace(tzinfo=TZ)


def turnback_segments(target):
    import pdfplumber
    out = []
    with pdfplumber.open(os.path.join(target, "route_turnback_notice_2027-03.pdf")) as pdf:
        text = "\n".join(p.extract_text() or "" for p in pdf.pages)
    for m in re.finditer(r"^SR \d+ (\d{7}) \d+\.\d{3} \d+\.\d{3}", text, re.M):
        out.append(int(m.group(1)))
    return out, text


def load(target):
    """The pack, parsed once."""
    from openpyxl import load_workbook
    p = {}
    p["inv"] = {int(r["segment_id"]): dict(sec=int(r["section"]), route=r["route"], bmp=float(r["begin_mp"]),
                                           emp=float(r["end_mp"]), length=float(r["length_mi"]),
                                           lanes=int(r["through_lanes"]), fc=r["functional_class"])
                for r in _csv(target, "road_inventory_2026.csv")}
    p["xw"] = _csv(target, "road_inventory_crosswalk_2025_2026.csv")
    p["turnback"], p["turnback_text"] = turnback_segments(target)
    p["det"] = _csv(target, "kerbline_detections_2027-03.csv")
    p["trav"] = _csv(target, "fleet_camera_traversals_2027-03.csv")
    p["wo"] = _csv(target, "mms_cold_patch_orders_2027-03.csv")
    p["covers"] = _csv(target, "utility_cover_inventory_2026.csv")
    wb = load_workbook(os.path.join(target, "kerbline_unit_register.xlsx"), read_only=True)
    rows = list(wb["Units"].iter_rows(values_only=True))
    p["reg"] = [dict(zip(rows[0], r)) for r in rows[1:]]
    return p


# ---------------------------------------------------------------------------------------------- shared pieces
def roadway(p, keep_turnback=False):
    """The program's roadway: the inventory's state segments in each section as at 31 March 2027."""
    tb = set() if keep_turnback else set(p["turnback"])
    return {s: v for s, v in p["inv"].items() if s not in tb}


def lane_miles(p, road):
    out = defaultdict(float)
    for s, v in road.items():
        out[v["sec"]] += v["length"] * v["lanes"]
    return dict(out)


def ref_to_2026(p, drop_unlisted=False, first_child=False):
    """road_ref (2025) and road_mp to the 2026 segment."""
    xw = defaultdict(list)
    for r in p["xw"]:
        xw[int(r["segment_id_2025"])].append((int(r["segment_id_2026"]), float(r["begin_mp_2026"]), float(r["end_mp_2026"])))

    def f(ref, mp):
        if ref == "":
            return None
        ref = int(ref)
        if ref in xw:
            if drop_unlisted:
                return None
            if first_child:
                return xw[ref][0][0]
            for seg, lo, hi in xw[ref]:
                if lo <= mp <= hi:
                    return seg
            return xw[ref][0][0]
        return ref if ref in p["inv"] else None
    return f


def units(p):
    """unit -> (vehicle, fleet group, installed date, removed date or None)."""
    out = {}
    for r in p["reg"]:
        inst = r["Installed"].date() if hasattr(r["Installed"], "date") else r["Installed"]
        rem = r["Removed"].date() if hasattr(r["Removed"], "date") and r["Removed"] is not None else r["Removed"]
        out[r["Unit serial"]] = (r["Vehicle"], r["Fleet group"], inst, rem)
    return out


def route_segment(p):
    """(route, milepost) -> the 2026 segment holding it."""
    by_route = defaultdict(list)
    for seg, v in p["inv"].items():
        by_route[v["route"]].append((v["bmp"], v["emp"], seg))
    for k in by_route:
        by_route[k].sort()

    def f(route, mp):
        for lo, hi, seg in by_route.get(route, ()):
            if lo - 1e-9 <= mp <= hi + 1e-9:
                return seg
        return None
    return f


def passes(p, flip=False):
    """Each detection's pass: its unit's vehicle (the register), its capture time in UTC (district local time read
    across the 14 March change), and the vehicle's traversal running at that time, give the pass's direction and
    speed (the 2026 segment's length over the traversal's duration). flip reads the direction convention the other
    way. Cached on the pack."""
    key = "_pass_flip" if flip else "_pass"
    if key in p:
        return p[key]
    u = units(p)
    by_veh = defaultdict(list)
    for t in p["trav"]:
        a, b = _utc(t["enter_utc"]), _utc(t["exit_utc"])
        seg = int(t["segment_id"])
        by_veh[t["vehicle_id"]].append((a, b, seg, t["direction"]))
    for k in by_veh:
        by_veh[k].sort()
    starts = {k: [x[0] for x in v] for k, v in by_veh.items()}
    slack = dt.timedelta(seconds=2.5)
    out = {}
    for d in p["det"]:
        if not d["road_mp"]:
            continue
        veh = u[d["unit_serial"]][0]
        t = _local(d["captured_local"]).astimezone(UTC)
        lst = by_veh.get(veh, [])
        i = bisect.bisect_right(starts.get(veh, []), t + slack) - 1
        hit = None
        for j in (i, i - 1):
            if 0 <= j < len(lst) and lst[j][0] - slack <= t <= lst[j][1] + slack:
                hit = lst[j]
                break
        if hit is None:
            continue
        a, b, seg, dirn = hit
        out[d["detection_id"]] = ((-1 if flip else 1) * UP[dirn], p["inv"][seg]["length"] / (b - a).total_seconds(),
                                  seg, a)
    p[key] = out
    return out


def pothole_detections(p, utc=False):
    """Kerbline's pothole-class detections on the state routes in March (district local time)."""
    out = []
    for d in p["det"]:
        if d["defect_class"] != "pothole" or not d["road_mp"]:
            continue
        t = _local(d["captured_local"])
        if utc:
            t = t.replace(tzinfo=UTC)
        if (t.year, t.month) != (2027, 3):
            continue
        out.append(d)
    return out


RADIUS = 0.05


def cover_errors(p, radius=RADIUS, flip=False):
    """unit -> [(signed milepost error in miles, direction sign, pass speed in miles per second)]: each utility_cover
    detection against the nearest surveyed cover on its route within radius."""
    ps = passes(p, flip)
    cov = defaultdict(list)
    for c in p["covers"]:
        cov[c["route"]].append(float(c["milepost"]))
    for k in cov:
        cov[k].sort()
    out = defaultdict(list)
    for d in p["det"]:
        if d["defect_class"] != "utility_cover" or not d["road_mp"] or d["detection_id"] not in ps:
            continue
        mp = float(d["road_mp"])
        xs = cov.get(d["route"], [])
        i = bisect.bisect_left(xs, mp)
        near = min((xs[j] for j in (i - 1, i) if 0 <= j < len(xs)), key=lambda c: abs(c - mp), default=None)
        if near is None or abs(near - mp) > radius:
            continue
        sg, v = ps[d["detection_id"]][:2]
        out[d["unit_serial"]].append((mp - near, sg, v))
    return out


def latency(p, how="ols", radius=RADIUS, flip=False):
    """Each unit's latency in seconds, from its cover detections: the error along the direction of travel against the
    pass's speed, through the origin (ols), as the median ratio (median), or as the mean of the two directions'
    through-the-origin fits (bydir)."""
    out = {}
    for u, xs in cover_errors(p, radius, flip).items():
        e = np.array([x[0] for x in xs])
        s = np.array([x[1] for x in xs])
        v = np.array([x[2] for x in xs])
        if how == "ols":
            out[u] = float(np.sum(e * s * v) / np.sum(v * v))
        elif how == "median":
            out[u] = float(np.median(e * s / v))
        else:
            fits = [float(np.sum(e[m] * s[m] * v[m]) / np.sum(v[m] * v[m])) for m in (s > 0, s < 0) if m.sum()]
            out[u] = float(np.mean(fits))
    return out


def positions(p, shift=None, dets=None):
    """[(route, milepost, detection)] for the March pothole detections; shift(d) is subtracted from road_mp."""
    dets = pothole_detections(p) if dets is None else dets
    out = []
    for d in dets:
        x = float(d["road_mp"])
        if shift is not None:
            k = shift(d)
            if k is None:
                continue
            x -= k
        out.append((d["route"], x, d))
    return out


def corrected(p, lat=None, speed="pass", dets=None, flip=False):
    """Each pothole detection moved back along its pass by its unit's latency times the pass's speed (speed='pass') or
    the unit's mean speed over its passes ('unit'); lat may be one latency for the fleet (a float)."""
    ps = passes(p, flip)
    lat = latency(p) if lat is None else lat
    if speed == "unit":
        per_unit = defaultdict(list)
        for d in p["det"]:
            if d["detection_id"] in ps:
                per_unit[d["unit_serial"]].append(ps[d["detection_id"]][1])
        vbar = {u: float(np.mean(v)) for u, v in per_unit.items()}

    def shift(d):
        if d["detection_id"] not in ps:
            return None
        sg, v = ps[d["detection_id"]][:2]
        L = lat if isinstance(lat, float) else lat[d["unit_serial"]]
        return sg * L * (vbar[d["unit_serial"]] if speed == "unit" else v)
    return positions(p, shift, dets)


def link(p, items, tol_ft, road, place="median", mapper=None):
    """Potholes per section: on each route the items sorted by milepost, an item joining the run before it when the
    gap is within tol_ft feet; a run is one pothole, placed on the 2026 segment holding its median (or mean) position,
    or (place='ref') on the 2025 road_ref of its median detection through the crosswalk."""
    seg_of = route_segment(p)
    mapper = mapper or ref_to_2026(p)
    by_route = defaultdict(list)
    for it in items:
        by_route[it[0]].append(it)
    out = defaultdict(int)
    for rt, xs in by_route.items():
        xs.sort(key=lambda t: (t[1], t[2]["detection_id"]))
        runs = [[xs[0]]]
        for a, b in zip(xs, xs[1:]):
            if (b[1] - a[1]) * FT > tol_ft:
                runs.append([b])
            else:
                runs[-1].append(b)
        for run in runs:
            ps_ = sorted(t[1] for t in run)
            mid = ps_[(len(ps_) - 1) // 2] if place != "mean" else float(np.mean(ps_))
            if place == "ref":
                d0 = min(run, key=lambda t: abs(t[1] - mid))[2]
                seg = mapper(d0["road_ref"], float(d0["road_mp"]))
            else:
                seg = seg_of(rt, mid)
            if seg in road:
                out[road[seg]["sec"]] += 1
    return dict(out)


def rates(p, counts, road):
    lm = lane_miles(p, road)
    return {sec: counts.get(sec, 0) / (lm[sec] / 10.0) for sec in lm}


def designated(r):
    return sorted(s for s, v in r.items() if v >= LINE)


TOL = 10.0


def golden(p):
    road = roadway(p)
    lat = latency(p)
    cnt = link(p, corrected(p, lat), TOL, road)
    est = rates(p, cnt, road)
    return dict(road=road, lat=lat, counts=cnt, est=est, set=designated(est))


def family(p, g=None):
    """Every licensed member: each unit's latency through the origin, as the median ratio or by direction, the covers
    matched within 0.05 or 0.1 mile, the corrected detections linked at 8 to 20 feet, a pothole placed at its median or
    its mean."""
    g = g or golden(p)
    out = {}
    for how in ("ols", "median", "bydir"):
        for radius in (0.05, 0.1):
            lat = latency(p, how, radius)
            items = corrected(p, lat)
            for tol in (8.0, 10.0, 12.0, 15.0, 20.0):
                for place in ("median", "mean"):
                    out["%s_r%.2f_%gft_%s" % (how, radius, tol, place)] = rates(p, link(p, items, tol, g["road"],
                                                                                       place=place), g["road"])
    return out


# ---------------------------------------------------------------------------------------------- the rungs
GRID = [0.5, 1, 2, 3, 4, 5, 6, 8, 10, 12, 15, 20, 25, 30, 35, 40, 45, 50, 55, 60, 70, 80, 90, 100, 110, 120, 130,
        140, 150, 160, 170, 180, 190, 200, 210, 230, 260, 300, 350, 400, 500, 600]


def single(p, tol_ft, road=None, items=None):
    """One linking distance for every detection as Kerbline placed it (the default path at any distance)."""
    road = road or roadway(p)
    return rates(p, link(p, positions(p) if items is None else items, tol_ft, road), road)


def fleet_latency(p):
    """One latency for the fleet: every unit's cover detections pooled into one fit through the origin."""
    e, s, v = [], [], []
    for xs in cover_errors(p).values():
        for x in xs:
            e.append(x[0]); s.append(x[1]); v.append(x[2])
    e, s, v = map(np.array, (e, s, v))
    return float(np.sum(e * s * v) / np.sum(v * v))


def unit_offsets(p):
    """Each unit's mean milepost error on the covers in each direction of travel, in miles (a fixed shift, no speed)."""
    out = {}
    for u, xs in cover_errors(p).items():
        for sg in (1, -1):
            es = [x[0] for x in xs if x[1] == sg]
            out[(u, sg)] = float(np.mean(es)) if es else 0.0
    return out


def offset_items(p):
    ps = passes(p)
    off = unit_offsets(p)

    def shift(d):
        if d["detection_id"] not in ps:
            return None
        return off[(d["unit_serial"], ps[d["detection_id"]][0])]
    return positions(p, shift)


def unit_rms(p):
    """Each unit's root-mean-square milepost error on the covers, in miles (the error read as scatter)."""
    return {u: float(np.sqrt(np.mean([x[0] ** 2 for x in xs]))) for u, xs in cover_errors(p).items()}


def rms_linkage(p, k=3.0, road=None):
    """Two detections joined within k times the root sum of squares of their units' RMS errors."""
    road = road or roadway(p)
    sig = unit_rms(p)
    seg_of = route_segment(p)
    by_route = defaultdict(list)
    for it in positions(p):
        by_route[it[0]].append(it)
    out = defaultdict(int)
    for rt, xs in by_route.items():
        xs.sort(key=lambda t: (t[1], t[2]["detection_id"]))
        runs = [[xs[0]]]
        for a, b in zip(xs, xs[1:]):
            tol = k * float(np.hypot(sig[a[2]["unit_serial"]], sig[b[2]["unit_serial"]]))
            if b[1] - a[1] > tol:
                runs.append([b])
            else:
                runs[-1].append(b)
        for run in runs:
            ps_ = sorted(t[1] for t in run)
            seg = seg_of(rt, ps_[(len(ps_) - 1) // 2])
            if seg in road:
                out[road[seg]["sec"]] += 1
    return rates(p, dict(out), road)


def max_per_pass(p, tol_ft, road=None):
    """Detections linked at one distance as placed, each run counted as the most detections one unit wrote on one pass
    inside it."""
    road = road or roadway(p)
    ps = passes(p)
    seg_of = route_segment(p)
    by_route = defaultdict(list)
    for it in positions(p):
        by_route[it[0]].append(it)
    out = defaultdict(int)
    for rt, xs in by_route.items():
        xs.sort(key=lambda t: (t[1], t[2]["detection_id"]))
        runs = [[xs[0]]]
        for a, b in zip(xs, xs[1:]):
            if (b[1] - a[1]) * FT > tol_ft:
                runs.append([b])
            else:
                runs[-1].append(b)
        for run in runs:
            per = defaultdict(int)
            for t in run:
                d = t[2]
                key = (d["unit_serial"], ps[d["detection_id"]][3] if d["detection_id"] in ps else d["captured_local"])
                per[key] += 1
            ps_ = sorted(t[1] for t in run)
            seg = seg_of(rt, ps_[(len(ps_) - 1) // 2])
            if seg in road:
                out[road[seg]["sec"]] += max(per.values())
    return rates(p, dict(out), road)


def mishandled(p):
    """The golden method with each Layer 1 item read the natural wrong way."""
    out = {}
    road = roadway(p)
    lat = latency(p)
    out["utc_window"] = rates(p, link(p, corrected(p, lat, dets=pothole_detections(p, utc=True)), TOL, road), road)
    items = corrected(p, lat)
    out["placed_by_road_ref"] = rates(p, link(p, items, TOL, road, place="ref"), road)
    out["crosswalk_unused"] = rates(p, link(p, items, TOL, road, place="ref", mapper=ref_to_2026(p, drop_unlisted=True)),
                                    road)
    out["crosswalk_first_child"] = rates(p, link(p, items, TOL, road, place="ref", mapper=ref_to_2026(p, first_child=True)),
                                         road)
    road_tb = roadway(p, keep_turnback=True)
    out["turnback_kept"] = rates(p, link(p, items, TOL, road_tb), road_tb)
    return out


# ---------------------------------------------------------------------------------------------- the ask
def _mms_time(s):
    return dt.datetime.strptime(s, "%m/%d/%Y %H:%M").replace(tzinfo=TZ) if s else None


def section_at(p):
    """at(route, mp): the section whose 2026 roadway holds the milepost."""
    segs = defaultdict(list)
    for v in p["inv"].values():
        segs[v["route"]].append((v["bmp"], v["emp"], v["sec"]))

    def at(route, mp):
        for lo, hi, sec in segs[route]:
            if lo <= mp <= hi:
                return sec
        return None
    return at


def ask_pounds(p, pounds="close", by="roadway"):
    """Ask A: per section, the pounds of cold mix its March cold patches used (COMPLETE orders finished in March), a
    pothole patched twice on each of its orders. pounds: 'close' takes once the figure a multi-line order's lines all
    carry with one close time (the data notes' order entry); 'per_order' every repeated figure once; 'lines' sums every
    line. by: 'roadway' places each line on the section whose roadway holds its milepost (program 3); 'order' takes
    the order's section column."""
    done = [r for r in p["wo"] if r["status"] == "COMPLETE" and r["finished"] and _mms_time(r["finished"]).month == 3]
    at = section_at(p)
    orders = defaultdict(list)
    for r in done:
        orders[r["wo_number"]].append(r)
    TOTALS.clear()
    lb = defaultdict(int)
    for wo, lines in orders.items():
        vals = [int(l["cold_mix_lb"]) for l in lines]
        repeated = len(lines) > 1 and len(set(vals)) == 1
        one_close = len({l["finished"] for l in lines}) == 1
        once = {"close": repeated and one_close, "per_order": repeated, "lines": False}[pounds]
        place = (lambda l: at(l["route"], float(l["milepost"]))) if by == "roadway" else (lambda l: int(l["section"]))
        if once:
            TOTALS.add(wo)
            lb[place(lines[0])] += vals[0]
        else:
            for l, v in zip(lines, vals):
                lb[place(l)] += v
    return dict(lb)


TOTALS = set()


def total_orders(p):
    """The orders whose repeated figure ask A's golden takes once."""
    ask_pounds(p)
    return set(TOTALS)


# ---------------------------------------------------------------------------------------------- distance sweeps
FINE = [round(0.5 + 0.25 * k, 2) for k in range(238)] + list(range(60, 602, 2))


def sweep_sets(p, items, road=None, grid=FINE):
    """distance in feet -> the designated set, for items linked at that one distance (numpy, for the fine grid)."""
    road = road or roadway(p)
    seg_of = route_segment(p)
    by = defaultdict(list)
    for it in items:
        by[it[0]].append(it[1])
    pre = {rt: np.sort(np.array(v)) for rt, v in by.items()}
    lm = lane_miles(p, road)
    out = {}
    for d in grid:
        cnt = defaultdict(int)
        for rt, xs in pre.items():
            cuts = np.where(np.diff(xs) * FT > d)[0]
            starts = np.concatenate(([0], cuts + 1))
            ends = np.concatenate((cuts + 1, [len(xs)]))
            for a, b in zip(starts, ends):
                seg = seg_of(rt, float(xs[a + (b - a - 1) // 2]))
                if seg in road:
                    cnt[road[seg]["sec"]] += 1
        out[d] = designated({s: cnt.get(s, 0) / (lm[s] / 10.0) for s in lm})
    return out


def bands(sets):
    """[(from, to, set)]: the runs of one designated set over a sweep."""
    out = []
    for d in sorted(sets):
        st = tuple(sets[d])
        if out and out[-1][2] == st:
            out[-1][1] = d
        else:
            out.append([d, d, st])
    return out
