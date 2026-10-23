"""The golden: every owed figure computed from the cut pack (bundle.py), written to generator/figures.json, and the three
deliverables, tbapcd_cerp_facilities_2027.xlsx, tbapcd_nox_inventory_2026.csv and tbapcd_facility_nox_2026.svg."""
import csv
import io
import json
import math
import os

import numpy as np

import bundle as B

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = "tbapcd_cerp_facilities_2027.xlsx"
CSV = "tbapcd_nox_inventory_2026.csv"
SVG = "tbapcd_facility_nox_2026.svg"
NUMBER = {n: no for n, no, _, _ in B.FACILITIES}


def compute(target):
    I = B.Inputs(target)
    k = I.k_pb1()
    tons = B.facility_tons(I, "golden", k)
    default = B.facility_tons(I, "latest")
    rank = sorted(B.ORDER, key=lambda n: -tons[n])
    five = B.five(tons)
    gas = B.gas_2026(target)
    eng = B.engine_lb(target)
    bl = B.read_billing(target)
    comb = B.combined_accounts(bl)
    el = B.elog_dates(target)
    outs = B.outages(target)
    rep = B.replacements(target)
    unit = B.unit_lb(I, "golden", k)
    pb_dev, pbm = I.pb1_months()
    burners = [d for d in I.open if "burner" in I.control[d]]
    f = {
        "call": {"facilities": five, "count": 5, "of": 14},
        "facilities": {n: {"number": NUMBER[n], "nox_tons": B.hround(tons[n]), "agreement": "Yes" if n in five else "No",
                           "gas_mmbtu": B.hround(gas[n]), "engine_nox_lb": B.hround(eng[n])} for n in B.ORDER},
        "components": {
            "rise_pct_per_pct": int(round(100 * (math.exp(k * 0.01) - 1))),
            "fifth": {"name": rank[4], "tons": B.hround(tons[rank[4]])},
            "sixth": {"name": rank[5], "tons": B.hround(tons[rank[5]])},
            "five_tons_low": min(B.hround(tons[n]) for n in five),
            "five_tons_high": max(B.hround(tons[n]) for n in five),
            "tests": int(len(I.tests)), "open_units": len(I.open), "monitored_units": len(I.monitored),
            "burner_units": len(burners), "devices": len(I.devices),
            "cems_tons": B.hround(sum(I.cems.values()) / 2000.0),
            "corridor_tons": B.hround(sum(tons.values())),
            "gas_total_mmbtu": B.hround(sum(gas.values())),
            "engine_total_lb": B.hround(sum(eng.values())),
        },
        "default": {"facilities": B.five(default), "tons": {n: B.hround(default[n]) for n in B.ORDER}},
        "ask_a": {"combined_accounts": sorted(comb), "lines": int(len(bl)), "first_therm_read": "July 2026"},
        "ask_b": {"elog": {n: str(el[n]) for n in B.ORDER if n in el}, "replaced": sorted(rep),
                  "bank_facilities": sorted(n for n in outs if n in ("Ashgrove Refining", "St. Aldric Medical Center",
                                                                     "Tamsett State University",
                                                                     "Quarry Lane Correctional Facility"))},
        "rank": rank,
        "stated": {"rule": "4306", "window": "calendar 2026", "agreements_by": "1 April 2027",
                   "engine_load_factor": 0.74, "g_lb": B.G_LB, "therms_per_mmbtu": 10, "plan_gas_section": "4.2",
                   "monitored_burner": pb_dev, "hmm_example": "0.30", "o2_ref_pct": 3, "heat_content_ref": B.HREF},
        "_unrounded": {"tons": {n: round(tons[n], 4) for n in B.ORDER}, "gas": {n: round(gas[n], 3) for n in B.ORDER},
                       "engine_lb": {n: round(eng[n], 3) for n in B.ORDER}, "k": k,
                       "unit_lb": {d: round(v, 2) for d, v in unit.items()},
                       "cems_lb": {d: round(v, 2) for d, v in I.cems.items()},
                       "pb1_months": {str(m): v for m, v in pbm.items()}},
    }
    return f


def write_csv(path, f):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(["facility_no", "facility", "rule_4306_nox_tons_2026"])
    for n in B.ORDER:
        w.writerow([f["facilities"][n]["number"], n, f["facilities"][n]["nox_tons"]])
    open(path, "w", encoding="utf-8", newline="").write(buf.getvalue())


def write_xlsx(path, f):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Agreements 2027"
    bold = Font(name="Calibri", size=11, bold=True)
    thin = Side(style="thin", color="808080")
    ws["A1"] = "NOx reduction agreements, 2027"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True)
    ws["A2"] = ("Compliance Division. Lowry Creek Community Emissions Reduction Plan 2027, section 2. Prepared "
                "19 February 2027.")
    ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="555555")
    ws.append([])
    ws.append(["Facility", "Facility no.", "2026 NOx, Rule 4306 units (tons): the agreement's baseline"])
    for c in ws[4]:
        c.font = bold
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
        c.border = Border(bottom=thin)
    for n in f["call"]["facilities"]:
        x = f["facilities"][n]
        ws.append([n, x["number"], x["nox_tons"]])
    for row in ws.iter_rows(min_row=5, max_row=9):
        row[1].number_format = "@"
        row[2].number_format = "#,##0"
    c6 = f["components"]["sixth"]
    ws["A11"] = ("In Appendix A order. Next in line: %s, %d tons. Agreements to be signed by 1 April 2027."
                 % (c6["name"], c6["tons"]))
    ws["A11"].font = Font(name="Calibri", size=9, color="555555")
    for col, wd in zip("ABC", (34, 11, 30)):
        ws.column_dimensions[col].width = wd
    ws.row_dimensions[4].height = 30
    ws.freeze_panes = "A5"

    s2 = wb.create_sheet("Facilities")
    s2.append(["Facility", "Facility no.", "2026 NOx, Rule 4306 units (tons)", "Agreement 2027",
               "Gas delivered by Tamsett Valley Gas, 2026 (MMBtu)", "Emergency standby engines' NOx, 2026 (lb)"])
    for c in s2[1]:
        c.font = bold
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
        c.border = Border(bottom=thin)
    fill = PatternFill("solid", fgColor="DDEBF7")
    for n in B.ORDER:
        x = f["facilities"][n]
        s2.append([n, x["number"], x["nox_tons"], x["agreement"], x["gas_mmbtu"], x["engine_nox_lb"]])
        if x["agreement"] == "Yes":
            for c in s2[s2.max_row]:
                c.fill = fill
    tot = s2.max_row + 1
    cp = f["components"]
    s2.append(["Corridor", None, cp["corridor_tons"], None, cp["gas_total_mmbtu"], cp["engine_total_lb"]])
    for c in s2[tot]:
        c.font = bold
        c.border = Border(top=thin)
    for row in s2.iter_rows(min_row=2, max_row=tot):
        row[1].number_format = "@"
        row[2].number_format = row[4].number_format = row[5].number_format = "#,##0"
        row[3].alignment = Alignment(horizontal="center")
    for col, wd in zip("ABCDEF", (34, 11, 14, 11, 18, 16)):
        s2.column_dimensions[col].width = wd
    s2.row_dimensions[1].height = 48
    s2.freeze_panes = "B2"
    s2.print_title_rows = "1:1"

    s3 = wb.create_sheet("Notes")
    notes = [
        ("2026 NOx, Rule 4306 units", "Plan 2.1 and 2.2: each facility's boilers, steam generators and process "
         "heaters under Rule 4306, calendar 2026. Units with a CEMS: the monitors' hourly NOx mass, substitute hours as "
         "reported. Units without one: each month's heat input from the emission statements at the rate the unit ran "
         "at on that month's gas. Their ultra-low and low-NOx burners set their combustion air by volume, so a richer "
         "gas fires hotter: the Lindqvist power boiler's monitor shows its NOx rate rising about %d per cent for each "
         "per cent rise in its BTU area's heat content (Tamsett Valley Gas's monthly table), and the %d source tests "
         "since 2018 bear the same rise out. Each burner's tests are carried from the heat content of their months' "
         "gas to each 2026 month's (Westside's moved from 1.027 to 1.087 therms a CCF in 2026, Eastside's held "
         "between 1.033 and 1.038)." % (cp["rise_pct_per_pct"], cp["tests"])),
        ("Gas delivered", "Plan 4.2: each billing period's therms on Tamsett Valley Gas's billing lines, spread evenly "
         "over the days from its prior read to the day before its current read, 10 therms to the MMBtu. Lines read "
         "before the utility's July 2026 billing conversion carry corrected CCF in the USAGE_THERMS column (times the "
         "line's BTU factor they are the therms billed); from the July reads the column is therms. The five "
         "combined-billing accounts carry the account's usage on each of their meters' lines, taken once a period."),
        ("Emergency engines' NOx", "Plan 4.3 and the inventory guidance section 6: hours x rated bhp x 0.74 x the "
         "tier's g/bhp-hr. Seven facilities moved their logs to the District's e-logbook during 2026, which records "
         "run time as hours and minutes (0.30 is half an hour); their engines' hour meters bear that reading out. At "
         "the four facilities whose engines run in parallel during an outage, each engine's emergency line carries the "
         "bank's engine-hours: each engine ran the outage the compliance log records. An engine replaced under the "
         "Clean Standby Engine program keeps its device number; its runs before the replacement date are the old "
         "engine's, at the old rating and tier the compliance log records."),
        ("Rounding", "Tons to the whole ton, gas to the whole MMBtu and engine NOx to the whole pound, each from the "
         "unrounded figure."),
        ("Sources", "CEMS hourly export 2026, source test databases 2018 to 2021 and 2022 to 2026, emission statements "
         "2018 to 2026, Tamsett Valley Gas monthly heat content, permit device list, compliance event log, Tamsett "
         "Valley Gas billing extract of 8 February 2027, Rule 4702 logs 2026 and engine inspections."),
    ]
    s3.append(["Item", "Note"])
    for c in s3[1]:
        c.font = bold
    for k, v in notes:
        s3.append([k, v])
    for row in s3.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[0].alignment = Alignment(vertical="top")
    s3.column_dimensions["A"].width = 26
    s3.column_dimensions["B"].width = 100
    wb.save(path)
    import postprocess as PP
    PP.normalize_ooxml(path, "Anna Haynes", "Andrew Greene", "2027-02-18T16:12:00Z", "2027-02-19T21:37:44Z",
                       excel=True, company=PP.COMPANY)


def write_svg(path, f):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    matplotlib.rcParams["svg.hashsalt"] = "task478-tbapcd"
    matplotlib.rcParams["svg.fonttype"] = "none"
    matplotlib.rcParams["font.family"] = "DejaVu Sans"
    names = sorted(B.ORDER, key=lambda n: -f["_unrounded"]["tons"][n])
    vals = [f["_unrounded"]["tons"][n] for n in names]
    on = [f["facilities"][n]["agreement"] == "Yes" for n in names]
    blue, grey, ink, muted = "#2a6fb0", "#c3c8ce", "#1f2328", "#59636e"
    fig, ax = plt.subplots(figsize=(9.6, 5.2), dpi=100)
    ys = np.arange(len(names))[::-1]
    ax.barh(ys, vals, height=0.66, color=[blue if o else grey for o in on], edgecolor="white", linewidth=2, zorder=3)
    for y, v, o in zip(ys, vals, on):
        ax.text(v + 0.25, y, "%d" % B.hround(v), ha="left", va="center", fontsize=9,
                color=ink if o else muted, fontweight="bold" if o else "normal")
    ax.set_yticks(ys)
    ax.set_yticklabels(names, fontsize=9, color=ink)
    ax.set_xlabel("NOx from Rule 4306 units, 2026 (tons)", fontsize=9, color=muted)
    ax.set_xlim(0, max(vals) * 1.12)
    ax.tick_params(axis="x", labelsize=8.5, colors=muted)
    ax.grid(axis="x", color="#e6e8eb", linewidth=0.8, zorder=0)
    for side in ("top", "right", "bottom"):
        ax.spines[side].set_visible(False)
    ax.spines["left"].set_color("#9aa1a9")
    from matplotlib.patches import Patch
    ax.legend(handles=[Patch(color=blue, label="NOx reduction agreement, 2027"), Patch(color=grey, label="No agreement")],
              loc="lower right", frameon=False, fontsize=8.5)
    c6 = f["components"]["sixth"]
    fig.suptitle("Five facilities' boilers and heaters emitted the most NOx in the Lowry Creek corridor in 2026", x=0.012,
                 ha="left", fontsize=11.5, fontweight="bold", color=ink)
    ax.set_title("Tons of NOx from each facility's Rule 4306 units; the five in blue get the plan's 2027 agreements.\n"
                 "Next in line: %s, %d tons." % (c6["name"], c6["tons"]), loc="left", fontsize=8.5, color=muted, pad=10,
                 linespacing=1.5)
    fig.text(0.012, 0.012, "Source: Tamsett Basin APCD 2026 point source inventory (CEMS, source tests, emission "
             "statements, Tamsett Valley Gas heat content).", fontsize=7.5, color=muted)
    fig.subplots_adjust(left=0.26, right=0.98, top=0.83, bottom=0.12)
    buf = io.StringIO()
    fig.savefig(buf, format="svg", metadata={"Date": None, "Creator": "Compliance Division"})
    plt.close(fig)
    svg = buf.getvalue().replace('id="matplotlib.', 'id="').replace("'DejaVu Sans'", "'DejaVu Sans', Arial, Helvetica, sans-serif")
    assert "matplotlib" not in svg.lower()
    open(path, "w", encoding="utf-8", newline="\n").write(svg)


def main(target=None):
    target = target or os.path.join(os.path.dirname(HERE), "target")
    f = compute(target)
    json.dump(f, open(os.path.join(HERE, "figures.json"), "w"), indent=1, sort_keys=True)
    out = os.path.join(os.path.dirname(HERE), "golden")
    os.makedirs(out, exist_ok=True)
    for fn in os.listdir(out):
        os.remove(os.path.join(out, fn))
    write_xlsx(os.path.join(out, XLSX), f)
    write_csv(os.path.join(out, CSV), f)
    write_svg(os.path.join(out, SVG), f)
    return f


if __name__ == "__main__":
    import sys
    f = main(sys.argv[1] if len(sys.argv) > 1 else None)
    print(json.dumps({k: v for k, v in f.items() if not k.startswith("_")}, indent=1)[:4000])
