"""Writes task478's evidence pack from world.py: the monitors' hourly export, the two source test databases, the
emission statements, the permit list, the compliance event log, the gas utility's billing extract, the facilities'
emergency engine logs and inspections, and the two distractors (the documents are docs.py's)."""
import csv
import datetime as dt
import os

import numpy as np

import world as W

FILES = {
    "plan": "tbapcd_lowry_creek_cerp_2027.pdf",
    "guidance": "tbapcd_emission_inventory_guidance_2026.pdf",
    "cems": "tbapcd_cems_hourly_2026.csv",
    "tests_old": "tbapcd_source_tests_2018_2021.csv",
    "tests_new": "tbapcd_source_tests_2022_2026.xlsx",
    "statements": "tbapcd_emission_statements_2017_2026.csv",
    "permits": "tbapcd_permit_devices.csv",
    "events": "tbapcd_compliance_events_2017_2026.csv",
    "billing": "tvg_commercial_billing_extract_2027-02.csv",
    "engine_logs": "tbapcd_rule4702_engine_logs_2026.xlsx",
    "inspections": "tbapcd_engine_inspections_2025_2027.csv",
    "no2": "tbapcd_lowry_creek_no2_hourly_2026.csv",
    "onroad": "tbapcd_lowry_creek_onroad_nox_2026.xlsx",
}
# the monitoring plan's unit ids for the monitored devices
CEMS_ID = {"D-1101": "CT1", "D-1102": "CT2", "D-1201": "H101", "D-1202": "H102", "D-1301": "MF1", "D-1701": "PB1",
           "D-1801": "GT3", "D-2101": "1", "D-2102": "2"}
ADDRESS = {"F01": "2200 Wexley Point Rd", "F02": "1 Refinery Ave", "F03": "880 Furnace St", "F04": "15500 Sedge Hollow Rd",
           "F05": "4917 Haskett Ridge Rd", "F06": "1300 Aldric Blvd", "F07": "60 Mill Race Ln", "F08": "500 University Dr",
           "F09": "7700 Briar Flats Rd", "F10": "210 Cannery Row", "F11": "3150 Harrow Ln", "F12": "4410 Orchard Way",
           "F13": "4486 Orchard Way", "F14": "9 Quarry Lane"}
# what each facility put its fuel in on its statements from 2021 (before that, every statement in MMSCF)
STATEMENT_UNIT = {"F01": "MMBTU", "F02": "THERMS", "F03": "MMSCF", "F04": "MMSCF", "F05": "MMSCF", "F06": "THERMS",
                  "F07": "MMBTU", "F08": "THERMS", "F09": "MMSCF", "F10": "MMSCF", "F11": "MMBTU", "F12": "THERMS",
                  "F13": "THERMS", "F14": "MMBTU"}
HHV_BIO = {"lfg": 534.0, "dg": 612.0}
EXTRACT_BILLING = dt.date(2027, 2, 3)
TESTERS = ("Basin Source Testing", "Delta Stack Services", "Westridge Emissions Measurement")


def _rng(k):
    return np.random.default_rng([W.SEED, 900 + k])


def fac_no(code, d=None):
    """The district's facility number on date d (Haskett Ridge was 0705 before its change of owner)."""
    if d is not None and code == W.RENUMBERED[0] and d < W.RENUMBERED[2]:
        return W.RENUMBERED[1]
    return W.FAC_NO[code]


def write_cems(path, e):
    rows = []
    for dev in W.MONITORED:
        u = W.UNIT[dev]
        for t, op, hi, r, sub in e["units"][dev]["cems"]:
            rows.append((W.FAC_NO[u["facility"]], CEMS_ID[dev], t.strftime("%Y-%m-%d"), t.hour, "%.2f" % op,
                         "%.1f" % hi, "%.4f" % r, "12" if sub else "01", "%.2f" % round(hi * r, 2)))
    rows.sort(key=lambda r: (r[0], r[1], r[2], r[3]))
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["facility_id", "unit_id", "op_date", "op_hour", "op_time", "heat_input_mmbtu", "nox_rate_lb_mmbtu",
                    "nox_rate_modc", "nox_mass_lb"])
        w.writerows(rows)
    return len(rows)


def _test_fields(dev, t, rng):
    """A test's reported quantities: heat input at the test load, ppm at the reference oxygen, lb/hr, fuel flow."""
    u = W.UNIT[dev]
    load = float(np.clip(u["load"] * rng.normal(1.0, 0.03), 0.6, 1.0))
    hi = u["rating"] * load
    fd = W.FD_NG if u["fuel"] == "ng" else W.FD_BIO
    o2ref = W.KINDS[u["kind"]][2]
    ppm = W.lb_to_ppm(t["measured"], o2ref, fd)
    hhv = W.HHV_NG if u["fuel"] == "ng" else HHV_BIO[u["fuel"]] * float(rng.uniform(0.97, 1.03))
    return dict(load=load, hi=hi, ppm=ppm, o2ref=o2ref, lbhr=t["measured"] * hi, scfh=hi * 1e6 / hhv, hhv=hhv)


def write_tests(old_path, new_path, e):
    import openpyxl
    from openpyxl.styles import Font, Alignment
    old, new = [], []
    tno = 18030
    for dev in sorted(W.OPEN):
        rng = _rng(int(dev[2:]))
        u = W.UNIT[dev]
        lim = W.KINDS[u["kind"]][1]
        for t in e["units"][dev]["tests"]:
            tno += int(rng.integers(5, 30))
            f = _test_fields(dev, t, rng)
            res = "PASS" if f["ppm"] <= lim else "FAIL"
            if t["date"].year <= 2021:
                old.append([str(tno), fac_no(u["facility"], t["date"]).lstrip("0"), dev[2:],
                            t["date"].strftime("%m/%d/%Y"), "%.2f" % f["ppm"], "%.0f" % f["o2ref"],
                            "%.3f" % f["lbhr"], "%.0f" % f["scfh"], "%.0f" % f["hhv"], "%.0f" % (100 * f["load"]), res])
            else:
                new.append(["ST-%d" % tno, fac_no(u["facility"], t["date"]), dev, t["date"], 3, round(f["ppm"], 2),
                            int(f["o2ref"]), round(t["measured"], 5), round(f["lbhr"], 3), round(f["hi"], 1),
                            round(100 * f["load"]), TESTERS[int(rng.integers(0, 3))], t["report"], res])
    old.sort(key=lambda r: dt.datetime.strptime(r[3], "%m/%d/%Y"))
    with open(old_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["TEST_NO", "FAC_NO", "DEVICE_NO", "TEST_DATE", "NOX_PPMVD", "O2_REF", "NOX_LB_HR", "FUEL_FLOW_SCFH",
                    "FUEL_HHV", "LOAD_PCT", "RESULT"])
        w.writerows(old)
    new.sort(key=lambda r: (r[3], r[2]))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    head = ["Test ID", "Facility No.", "Device No.", "Test Date", "Runs", "NOx (ppmvd @ ref O2)", "Ref O2 (%)",
            "NOx (lb/MMBtu)", "NOx (lb/hr)", "Heat Input (MMBtu/hr)", "Load (%)", "Source Tester", "Report Received",
            "Result"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in new:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        row[3].number_format = "yyyy-mm-dd"
        row[12].number_format = "yyyy-mm-dd"
        row[7].number_format = "0.00000"
    widths = [11, 10, 10, 11, 6, 12, 8, 12, 10, 11, 8, 30, 12, 8]
    for i, wd in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = wd
    ws.freeze_panes = "A2"
    ns = wb.create_sheet("Field notes")
    for line in (["Source test results database (new system, from January 2022)"],
                 ["Results are the average of the test's runs (normally three), corrected to the reference oxygen of"
                  " the device's permit limit."],
                 ["lb/MMBtu is computed by the tester under EPA Method 19 from the fuel's F-factor; lb/hr from the"
                  " measured stack flow."],
                 ["Tests before 2022 are in the earlier database's export."]):
        ns.append(line)
    ns.column_dimensions["A"].width = 110
    wb.save(new_path)
    return len(old), len(new)


def write_statements(path, e):
    rows = []
    for dev in sorted(W.UNIT):
        u = W.UNIT[dev]
        hrs = e["units"][dev]["hours"] if not u["monitored"] else W.run_hours(dev)
        hr = W.heat_rate(dev)
        rng = _rng(100 + int(dev[2:]))
        for y in range(2017, 2027):
            for m in range(1, 13):
                days = [d for d in hrs if d.year == y and d.month == m]
                h = sum(hrs[d] for d in days)
                if u["monitored"] and y == W.YEAR:
                    h = sum(op for t, op, _, _, _ in e["units"][dev]["cems"] if t.month == m)
                    mmbtu = sum(hi for t, _, hi, _, _ in e["units"][dev]["cems"] if t.month == m)
                else:
                    mmbtu = h * hr
                unit = "MMSCF" if y <= 2020 else STATEMENT_UNIT[u["facility"]]
                hhv = ""
                if u["fuel"] != "ng":
                    unit = "MMSCF"
                    hhv = "%.0f" % (HHV_BIO[u["fuel"]] * float(rng.uniform(0.97, 1.03)))
                    qty = "%.4f" % (mmbtu / float(hhv))
                elif unit == "MMSCF":
                    qty = "%.4f" % (mmbtu / W.HHV_NG)
                elif unit == "THERMS":
                    qty = "%.0f" % (mmbtu * 10)
                else:
                    qty = "%.1f" % mmbtu
                fuel = {"ng": "NATURAL GAS", "lfg": "LANDFILL GAS", "dg": "DIGESTER GAS"}[u["fuel"]]
                rows.append([str(y), fac_no(u["facility"], dt.date(y, m, 1)), dev, str(m), fuel, qty, unit, hhv,
                             "%.1f" % h])
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["STMT_YEAR", "FAC_NO", "DEVICE_NO", "MONTH", "FUEL", "QUANTITY", "UNIT", "HHV_BTU_SCF", "HOURS"])
        w.writerows(rows)
    return len(rows)


# devices the permit list carries beyond the plan's units and the standby engines (small units under Rule 4307)
SMALL = [("F02", "D-1290", "process water heater", 4.2), ("F03", "D-1390", "annealing lehr burner", 3.6),
         ("F06", "D-1690", "kitchen water heater", 2.4), ("F07", "D-1790", "space heater bank", 4.8),
         ("F08", "D-1890", "laboratory boiler", 3.0), ("F10", "D-2090", "can washer heater", 2.0),
         ("F12", "D-2290", "office boiler", 1.2), ("F13", "D-2390", "curing press boiler", 4.6),
         ("F14", "D-2490", "space heating boiler", 2.8)]
REISSUED = {"D-1102": dt.date(2026, 3, 30), "D-1801": dt.date(2026, 3, 30), "D-1701": dt.date(2026, 6, 15),
            "D-3202": dt.date(2026, 6, 1), "D-4402": dt.date(2026, 9, 21), "D-1290": dt.date(2026, 2, 9)}


def write_permits(path):
    rows = []
    for dev in sorted(W.UNIT):
        u = W.UNIT[dev]
        k = W.KINDS[u["kind"]]
        issued = W.RENUMBERED[2] if u["facility"] == W.RENUMBERED[0] else (u["catalysts"][0] if u["catalysts"] else dt.date(2017, 9, 1))
        issued = REISSUED.get(dev, issued)
        rows.append([W.FAC_NO[u["facility"]], W.FAC_NAME[u["facility"]], dev, u["desc"], "4306",
                     "%g" % u["rating"], "MMBtu/hr", {"ng": "natural gas", "lfg": "landfill gas", "dg": "digester gas"}[u["fuel"]],
                     "SCR", "CEMS" if u["monitored"] else "source test", CEMS_ID.get(dev, ""),
                     "" if u["kind"] == "furnace" else "%g" % k[1], "" if u["kind"] == "furnace" else "%g" % k[2],
                     "" if u["monitored"] else str(12 * u["cycle"]), "", "", issued.isoformat()])
    for dev, e in sorted(W.ENGINE.items()):
        issued = e["replaced"][0] if e["replaced"] else e["installed"]
        issued = REISSUED.get(dev, issued)
        rows.append([W.FAC_NO[e["facility"]], W.FAC_NAME[e["facility"]], dev, "emergency standby diesel engine", "4702",
                     str(e["bhp"]), "bhp", "diesel", "", "", "", "", "", "", e["model"], e["tier"], issued.isoformat()])
    for code, dev, desc, rating in SMALL:
        rows.append([W.FAC_NO[code], W.FAC_NAME[code], dev, desc, "4307", "%g" % rating, "MMBtu/hr", "natural gas",
                     "low-NOx burner", "tune-up", "", "", "", "", "", "", REISSUED.get(dev, dt.date(2019, 5, 6)).isoformat()])
    rows.sort(key=lambda r: (r[0], r[2]))
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["FAC_NO", "FACILITY", "DEVICE_NO", "DESCRIPTION", "RULE", "RATING", "RATING_UNITS", "FUEL",
                    "NOX_CONTROL", "MONITORING", "CEMS_UNIT_ID", "NOX_LIMIT_PPMVD", "REF_O2_PCT", "TEST_FREQ_MO",
                    "ENGINE_MODEL", "ENGINE_TIER", "PERMIT_ISSUED"])
        w.writerows(rows)
    return len(rows)


CAT_TEXT = ("SCR catalyst replaced, two layers (vendor changeout)", "Catalyst modules replaced in SCR reactor",
            "SCR catalyst changeout completed; unit returned to service", "New SCR catalyst installed",
            "Replaced spent SCR catalyst")


def events(e):
    """The compliance event log's rows: (date, facility code, device, type, description)."""
    out = []
    for dev, u in W.UNIT.items():
        rng = _rng(300 + int(dev[2:]))
        cats = list(W.schedule(dev)[0])
        for i, c in enumerate(cats):
            if i == 0:
                out.append((c, u["facility"], dev, "SCR INSTALLED", "SCR system installed and started up (Rule 4306"
                            " retrofit, compliance date 1 January 2018)"))
            else:
                out.append((c, u["facility"], dev, "CATALYST", CAT_TEXT[int(rng.integers(0, len(CAT_TEXT)))]))
        for t in e["units"][dev].get("tests", []):
            if t["kind"] == "retest":
                out.append((t["date"] - dt.timedelta(days=int(rng.integers(10, 21))), u["facility"], dev,
                            "SOURCE TEST", "Retest notification received; test scheduled %s"
                            % t["date"].strftime("%m/%d/%Y")))
            else:
                out.append((t["date"] - dt.timedelta(days=int(rng.integers(30, 45))), u["facility"], dev,
                            "SOURCE TEST", "Source test notification received; test scheduled %s"
                            % t["date"].strftime("%m/%d/%Y")))
            if t["nov"] and t["nov"] <= W.END:
                out.append((t["nov"], u["facility"], dev, "NOV", "Notice of violation: NOx at the source test of %s "
                            "above the permit limit; corrective action and a retest required"
                            % t["date"].strftime("%m/%d/%Y")))
        for ch in u["calendar"]:
            if ch[0] == "defer":
                out.append((ch[1] - dt.timedelta(days=9), u["facility"], dev, "TEST DEFERRAL",
                            "Source test due %s deferred to %s: unit out of service for a scheduled engine overhaul"
                            % (ch[1].strftime("%m/%d/%Y"), ch[2].strftime("%m/%d/%Y"))))
        if dev in W.DOWN:
            a, b = W.DOWN[dev]
            out.append((a, u["facility"], dev, "BREAKDOWN", "Unit shut down for scheduled overhaul; returned to service"
                        " %s" % b.strftime("%m/%d/%Y")))
        for y in range(2017, 2027):
            if rng.random() < 0.45:
                d = dt.date(y, int(rng.integers(1, 13)), int(rng.integers(1, 28)))
                out.append((d, u["facility"], dev, "BREAKDOWN", "Breakdown reported: %s; corrected within 24 hours"
                            % ("ammonia injection skid trip", "burner management fault", "SCR inlet temperature low",
                               "control system communication loss")[int(rng.integers(0, 4))]))
            if rng.random() < 0.6:
                d = dt.date(y, int(rng.integers(1, 13)), int(rng.integers(1, 28)))
                out.append((d, u["facility"], dev, "INSPECTION", "Annual compliance inspection; no violations noted"))
        if u["monitored"]:
            for y in range(2017, 2027):
                out.append((dt.date(y, 4 + int(rng.integers(0, 4)), int(rng.integers(1, 28))), u["facility"], dev,
                            "RATA", "CEMS relative accuracy test audit passed"))
    for dev, en in W.ENGINE.items():
        rng = _rng(500 + int(dev[2:]))
        if en["replaced"]:
            d, bhp, tier, model = en["replaced"]
            out.append((d, en["facility"], dev, "EQUIPMENT", "Emergency engine replaced under the district's Clean"
                        " Standby Engine program: removed %s (%s bhp, %s); installed %s (%s bhp, %s)"
                        % (model, "{:,}".format(bhp), tier, en["model"], "{:,}".format(en["bhp"]), en["tier"])))
        if rng.random() < 0.5:
            out.append((dt.date(2026, int(rng.integers(1, 13)), int(rng.integers(1, 28))), en["facility"], dev,
                        "INSPECTION", "Emergency engine inspection; hour meter and log reviewed"))
    code, old, d, prev, new = W.RENUMBERED
    out.append((d, code, "", "OWNERSHIP", "Change of ownership: %s to %s; facility number %s replaced by %s"
                % (prev, new, old, W.FAC_NO[code])))
    out.append((dt.date(2026, 3, 30), "F01", "D-1102", "PERMIT MOD", "Permit reissued: administrative amendment"
                " (monitoring plan update)"))
    return sorted((r for r in out if r[0] <= W.END), key=lambda r: (r[0], r[2], r[3]))


def write_events(path, e):
    rows = events(e)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["EVENT_ID", "FAC_NO", "DEVICE_NO", "EVENT_DATE", "EVENT_TYPE", "DESCRIPTION"])
        for i, (d, code, dev, typ, desc) in enumerate(rows):
            date = d.strftime("%m/%d/%Y") if d.year <= 2020 else d.isoformat()
            w.writerow(["E%06d" % (100000 + 7 * i + 3), fac_no(code, d), dev, date, typ, desc])
    return len(rows)


def write_billing(path, e):
    rows = []
    for b in e["bills"]:
        if b["read_to"] + dt.timedelta(days=3) > EXTRACT_BILLING:
            continue                    # billed after the utility pulled the extract
        code = b["facility"]
        rows.append([str(b["bill"]), b["account"], ADDRESS[code].upper().replace(" RD", " RD").replace("WAY", "WY"),
                     b["meter"], b["read_from"].strftime("%m/%d/%Y"), b["read_to"].strftime("%m/%d/%Y"),
                     str(b["prior"]), str(b["current"]), "A", "" if b["evc"] else "%.3f" % b["mult"],
                     "%.3f" % b["btu"], str(b["usage"]),
                     "G-20" if b["mult"] > 1 else "G-10",
                     (b["read_to"] + dt.timedelta(days=3)).strftime("%m/%d/%Y")])
    rows.sort(key=lambda r: (r[1], r[3], dt.datetime.strptime(r[4], "%m/%d/%Y")))
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["BILL_ID", "ACCOUNT_NO", "SERVICE_ADDRESS", "METER_NO", "READ_FROM", "READ_TO", "PRIOR_READ",
                    "CURRENT_READ", "READ_TYPE", "PRESSURE_MULT", "BTU_FACTOR", "USAGE_CCF", "RATE_CODE", "BILL_DATE"])
        w.writerows(rows)
    return len(rows)


def _hmm(h):
    """Hours as the paper logs write them: whole hours, a point, minutes."""
    mins = int(round(h * 60))
    return "%d.%02d" % (mins // 60, mins % 60)


def write_engine_logs(path, e):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    n = 0
    for code in sorted(W.FAC_NAME, key=lambda c: W.FAC_NO[c]):
        devs = [d for d, en in W.ENGINE.items() if en["facility"] == code]
        if not devs:
            continue
        ws = wb.create_sheet(W.FAC_NO[code])
        ws.append(["%s (facility %s): Rule 4702 emergency engine operating log, calendar 2026"
                   % (W.FAC_NAME[code], W.FAC_NO[code])])
        ws["A1"].font = Font(bold=True)
        if code in W.HMM:
            ws.append(["Date", "Engine (device no.)", "Hours", "Reason"])
        else:
            ws.append(["Date", "Device", "Run hours", "Purpose"])
        for c in ws[2]:
            c.font = Font(bold=True)
        rows = []
        for dev in devs:
            for d, h, why in e["runs"][dev]:
                v = _hmm(h) if code in W.HMM else "%.1f" % h
                rows.append([d, dev, float(v), {"readiness": "Testing (weekly)", "load test": "Maintenance (load bank)",
                                                "emergency": "Emergency (utility outage)"}[why]])
        rows.sort(key=lambda r: (r[0], r[1]))
        for r in rows:
            ws.append(r)
            n += 1
        for row in ws.iter_rows(min_row=3):
            row[0].number_format = "mm/dd/yyyy"
            row[2].number_format = "0.00" if code in W.HMM else "0.0"
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 10
        ws.column_dimensions["D"].width = 28
    wb.save(path)
    return n


def write_inspections(path, e):
    rows = []
    insp = ("J. Okafor", "M. Delgado", "S. Whitfield")
    for i, (dev, d, read) in enumerate(sorted(e["inspections"], key=lambda r: (r[1], r[0]))):
        rows.append([dev, W.FAC_NO[W.ENGINE[dev]["facility"]], d.isoformat(), "%.1f" % read, insp[i % 3], "compliant"])
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["DEVICE_NO", "FAC_NO", "INSPECTION_DATE", "HOUR_METER", "INSPECTOR", "RESULT"])
        w.writerows(rows)
    return len(rows)


def write_no2(path):
    rng = _rng(77)
    sites = (("LC-01", "Lowry Creek Elementary"), ("LC-02", "Harrow Lane Park"), ("LC-03", "Seddon Flats Library"))
    rows = []
    for code, _ in sites:
        base = float(rng.uniform(9, 14))
        for d in W.DAYS:
            if d.year != W.YEAR:
                continue
            for h in range(24):
                v = base * (1 + .5 * np.cos(2 * np.pi * (h - 7) / 24) + .3 * np.cos(2 * np.pi * (d.timetuple().tm_yday - 15) / 365))
                v *= float(np.exp(rng.normal(0, .25)))
                q = "" if rng.random() > 0.008 else "AN"
                rows.append([code, d.isoformat(), str(h), "" if q else "%.1f" % v, q])
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["site_id", "date_lst", "hour_lst", "no2_ppb", "qualifier"])
        w.writerows(rows)
    return len(rows)


def write_onroad(path):
    import openpyxl
    from openpyxl.styles import Font
    rng = _rng(78)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Links 2026"
    ws.append(["Link ID", "Road", "From", "To", "Length (mi)", "AADT", "Heavy-duty share", "VMT (million)",
               "NOx (tons)"])
    for c in ws[1]:
        c.font = Font(bold=True)
    roads = ("State Route 41", "Lowry Creek Pkwy", "Harrow Ln", "Orchard Way", "Refinery Ave", "Seddon Flats Rd",
             "Interstate 7")
    k = 0
    for road in roads:
        for j in range(int(rng.integers(3, 7))):
            k += 1
            ln = round(float(rng.uniform(0.4, 2.6)), 2)
            aadt = int(rng.integers(4000, 61000))
            hd = round(float(rng.uniform(0.03, 0.18)), 3)
            vmt = ln * aadt * 365 / 1e6
            nox = vmt * (0.18 + 2.1 * hd) / 2000 * 1e6 / 453.6 * 0.4536
            ws.append(["L%03d" % k, road, "node %d" % (k * 2), "node %d" % (k * 2 + 1), ln, aadt, hd, round(vmt, 2),
                       round(nox / 1000, 2)])
    for col, wd in zip("ABCDEFGHI", (9, 18, 10, 10, 10, 9, 14, 13, 10)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A2"
    wb.save(path)
    return k


def build_pack(out, e=None):
    e = e or W.everything()
    f = lambda k: os.path.join(out, FILES[k])
    n = {}
    n["cems"] = write_cems(f("cems"), e)
    n["tests_old"], n["tests_new"] = write_tests(f("tests_old"), f("tests_new"), e)
    n["statements"] = write_statements(f("statements"), e)
    n["permits"] = write_permits(f("permits"))
    n["events"] = write_events(f("events"), e)
    n["billing"] = write_billing(f("billing"), e)
    n["engine_logs"] = write_engine_logs(f("engine_logs"), e)
    n["inspections"] = write_inspections(f("inspections"), e)
    n["no2"] = write_no2(f("no2"))
    n["onroad"] = write_onroad(f("onroad"))
    return e, n
