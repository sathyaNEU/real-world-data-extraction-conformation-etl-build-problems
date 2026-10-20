"""The pack's documents and workbooks: Article 14 of the Municipal Fund Regulation 2027, the Bureau's report on the
Register Coverage Audit 2026, the register's record layouts, the ministry's delivery specification, the municipality
code workbook and the 2024 pilot's workbook."""
import datetime as dt

import docx
import openpyxl
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import A4                                                # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import cm                                                    # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import world as W    # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("LiberationSerif", "Regular"), ("LiberationSerif-Bold", "Bold"), ("LiberationSerif-Italic", "Italic"),
              ("LiberationSans", "Regular"), ("LiberationSans-Bold", "Bold"), ("LiberationSans-Italic", "Italic")):
    try:
        pdfmetrics.getFont(nm)
    except KeyError:
        pdfmetrics.registerFont(TTFont(nm, FONT_DIR + nm.split("-")[0] + "-" + f + ".ttf"))
INK = colors.HexColor("#1b1d1f")
GREY = colors.HexColor("#5b6067")
NAVY = colors.HexColor("#1f3a5f")
OCHRE = colors.HexColor("#7a5a12")
SANS = ParagraphStyle("sans", fontName="LiberationSans", fontSize=9.4, leading=13.0, textColor=INK, spaceAfter=6)
SERIF = ParagraphStyle("serif", fontName="LiberationSerif", fontSize=10.6, leading=14.6, textColor=INK, spaceAfter=7)
CELL = ParagraphStyle("cell", fontName="LiberationSans", fontSize=8.0, leading=10.0, textColor=INK)
CELLB = ParagraphStyle("cellb", parent=CELL, fontName="LiberationSans-Bold")
SMALL = ParagraphStyle("small", fontName="LiberationSans", fontSize=7.8, leading=10.2, textColor=GREY)
HEAD_FILL = PatternFill("solid", fgColor="DCE3EA")


def styles(accent, serif=False):
    face = "LiberationSerif" if serif else "LiberationSans"
    h1 = ParagraphStyle("h1", fontName=face + "-Bold", fontSize=14, leading=18, textColor=accent, spaceAfter=2)
    h1s = ParagraphStyle("h1s", fontName=face, fontSize=9.2, leading=12.5, textColor=GREY, spaceAfter=10)
    h2 = ParagraphStyle("h2", fontName=face + "-Bold", fontSize=10.4, leading=14, textColor=INK, spaceBefore=7,
                        spaceAfter=3, keepWithNext=1)
    return h1, h1s, h2


def _doc(path, title, author, footer, serif=False):
    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(GREY)
        canvas.setFont("LiberationSerif" if serif else "LiberationSans", 7.6)
        canvas.drawString(2.2 * cm, 1.4 * cm, footer)
        canvas.drawRightString(A4[0] - 2.2 * cm, 1.4 * cm, "%d" % doc.page)
        canvas.restoreState()
    d = SimpleDocTemplate(path, pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=2.0 * cm,
                          bottomMargin=2.2 * cm, title=title, author=author, subject="", creator="", producer="")
    return d, frame


def _table(rows, widths, head="#e3e9ef", right=()):
    data = []
    for i, r in enumerate(rows):
        line = []
        for j, v in enumerate(r):
            st = CELLB if i == 0 else CELL
            if i > 0 and j in right:
                st = ParagraphStyle("r", parent=st, alignment=2)
            line.append(Paragraph(str(v), st))
        data.append(line)
    t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa0a6")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(head)),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 4),
                           ("RIGHTPADDING", (0, 0), (-1, -1), 4)]))
    return t


def _n(v):
    return "{:,}".format(int(v))


# ----------------------------------------------------------------------------------------------- the regulation
def write_regulation(path):
    h1, h1s, h2 = styles(NAVY, serif=True)
    d, frame = _doc(path, "Municipal Fund Regulation 2027, Article 14", W.MINISTRY,
                    "Municipal Fund Regulation 2027  |  Official Gazette 2026, no. 31,204  |  30 June 2026", serif=True)
    s = [Paragraph("Municipal Fund Regulation 2027", h1),
         Paragraph("Regulation of the Minister of the Interior of 30 June 2026 on the distribution of the Municipal "
                   "Fund for 2027. Extract: Chapter 4, Article 14, with its explanatory note.", h1s),
         Paragraph("Article 14. Small-municipality supplement", h2),
         Paragraph("1. For 2027 a municipality receives the small-municipality supplement if its adjusted resident "
                   "population on 1 January 2026 is below 5,000.", SERIF),
         Paragraph("2. The adjusted resident population of a municipality is the number of persons registered as "
                   "residents of the municipality in the Population Register at the start of 1 January 2026, less the "
                   "number of those persons estimated from the %s's %s not to have been living in the municipality "
                   "on that date." % (W.BUREAU, W.AUDIT), SERIF),
         Paragraph("3. The Bureau delivers each municipality's adjusted resident population to the Minister rounded to "
                   "the nearest whole person, and the rounded figure is compared with 5,000.", SERIF),
         Paragraph("4. Municipalities are taken as constituted on 1 January 2026.", SERIF),
         Paragraph("5. The supplement for 2027 is EUR 418,000 for each municipality that receives it. It is paid in "
                   "twelve monthly instalments with the general allowance.", SERIF),
         Paragraph("Explanatory note to Article 14", h2),
         Paragraph("Until 2026 the supplement followed the register count. A person who leaves the country stays on "
                   "the register until somebody reports the departure or the municipality ends the registration after "
                   "an address investigation, so the register count includes people who no longer live in the "
                   "municipality. In 2026 the Bureau audited, for the first time, a sample of the persons on the "
                   "register to measure how many were not living in their registered municipality. From 2027 the "
                   "supplement follows the register count less the persons the audit shows were not living there.", SERIF),
         Paragraph("The Bureau publishes the method by which it applies the audit to municipalities in its report on "
                   "the audit. Paragraph 4 means that the municipality of Ravenmeer, formed on 1 January 2026 from "
                   "Ravensloot and Meerdijk, is one municipality for 2027.", SERIF)]
    d.build(s, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------------------------------------- the audit report
def write_report(path, national):
    """national: dict result code -> sampled persons with that final result"""
    h1, h1s, h2 = styles(OCHRE)
    d, frame = _doc(path, "Register Coverage Audit 2026: design, fieldwork and national results", W.PEOPLE["method"],
                    "%s  |  %s  |  Register Coverage Audit 2026  |  September 2026" % (W.BUREAU, W.DIVISION))
    total = sum(national.values())
    absent = sum(national[c] for c in ("20", "30", "40"))
    est = absent / total * W.NATIONAL_REGISTER
    labels = {"10": "living at the registered address", "11": "living elsewhere in the registered municipality",
              "50": "living in the municipality on 1 January 2026, left afterwards",
              "20": "living elsewhere in the country, not registered there",
              "30": "living abroad", "40": "died before 1 January 2026, death not recorded"}
    rows = [["Final result", "Code", "Sampled persons", "Share"]]
    for c in ("10", "11", "50", "20", "30", "40"):
        rows.append([labels[c], c, _n(national[c]), "%.2f%%" % (100.0 * national[c] / total)])
    rows.append(["All sampled persons", "", _n(total), "100.00%"])
    s = [Paragraph("Register Coverage Audit 2026", h1),
         Paragraph("Design, fieldwork and national results. %s, %s. Prepared by %s, methodologist; approved by %s, "
                   "head of population statistics. Published 18 September 2026." % (
                       W.BUREAU, W.DIVISION, W.PEOPLE["method"], W.PEOPLE["head"]), h1s),
         Paragraph("1. Purpose", h2),
         Paragraph("The audit measures how many persons registered in the Population Register at the start of 1 "
                   "January 2026 were not living in their registered municipality on that date. It is the audit "
                   "Article 14 of the Municipal Fund Regulation 2027 names for the adjusted resident population.", SANS),
         Paragraph("2. Design", h2),
         Paragraph("A simple random sample of %s persons was drawn from the %s persons registered in the Population "
                   "Register at the start of 1 January 2026, every registered person with the same chance of being "
                   "drawn. For each sampled person the Bureau extracted the register record as it stood that day and "
                   "the person's register history; field staff used both to trace the person." % (
                       _n(W.N_AUDIT), _n(W.NATIONAL_REGISTER)), SANS),
         Paragraph("3. Fieldwork", h2),
         Paragraph("Between 2 February and 29 May 2026 field staff visited each sampled person's registered address. "
                   "Where nobody could be found they called again, spoke to neighbours and the person's household and "
                   "checked other administrative records (benefits, tax, health insurance and, through the European "
                   "exchange, registration in another member state). Every case was resolved to a final result "
                   "describing where the person was living on 1 January 2026; a visit without contact is recorded "
                   "as 91 and is not a final result. Result 50 marks persons who were living in the municipality on "
                   "1 January 2026 and left it before the visit: they count as living there.", SANS),
         Paragraph("4. National results", h2),
         _table(rows, [8.2 * cm, 1.4 * cm, 3.0 * cm, 2.4 * cm], right=(2, 3)),
         Spacer(1, 5),
         Paragraph("Results 20, 30 and 40 are persons not living in their registered municipality on 1 January 2026: "
                   "%s of the %s sampled, %.2f per cent. Applied to the register, about %s registered persons were "
                   "not living in their registered municipality." % (
                       _n(absent), _n(total), 100.0 * absent / total, _n(round(est, -3))), SANS),
         Paragraph("5. Figures for municipalities", h2),
         Paragraph("The audit sampled about thirteen of every ten thousand registered persons, too few in any one "
                   "municipality for that municipality's own sampled persons to give its figure. A municipality's "
                   "figure is reached by giving each person on its register the share not living in the "
                   "municipality that the audit found among sampled persons like that person in the respects that "
                   "bear on it, and adding those shares.", SANS),
         Paragraph("6. Files", h2),
         Paragraph("The sample (one record per sampled person, as the register stood at the start of 1 January "
                   "2026), the fieldwork visits and the sampled persons' register histories are supplied with this "
                   "report in the layouts of the Population Register's statistical extracts.", SANS)]
    d.build(s, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------------------------------------- record layouts
LAYOUTS = [
    ("Population Register extract, 1 January 2026 (population_register_extract_20260101.txt)",
     "Persons registered as residents of the in-play municipalities at the start of 1 January 2026; pipe-delimited, "
     "a header line and a trailer line (TRAILER, record count, extract date).",
     [("PERSON_KEY", "pseudonymised person key, the same in every extract"),
      ("MUN_CODE", "municipality of registration, four digits"),
      ("ADDRESS_ID", "address identifier, sixteen digits; the first four are the code of the municipality that "
                     "issued it"),
      ("ADDRESS_SINCE", "date from which the person is registered at the current address (YYYYMMDD)"),
      ("SEX", "M or F"), ("BIRTH_DATE", "YYYYMMDD; a day or month that is not known is written 00"),
      ("BIRTH_COUNTRY", "country of birth, country code"), ("NATIONALITY", "nationality, country code"),
      ("MARITAL_STATUS", "O never married or in a registered partnership, H married or in a registered partnership")]),
    ("Register events (population_register_events_extract.csv, register_coverage_audit_2026_histories.csv)",
     "Every event the register holds for each person, in sequence.",
     [("PERSON_KEY", "person key"), ("SEQ", "sequence number of the event for the person"),
      ("EVENT", "event code (list below)"), ("EVENT_DATE", "date the event took effect (YYYYMMDD)"),
      ("MUN_CODE", "municipality where the event was registered, its code at the time"),
      ("COUNTRY", "IMM: country the person came from; EMI: country of destination"),
      ("FROM_MUN_CODE", "VER: municipality the person moved from"),
      ("ANNULS_SEQ", "ANN: the SEQ of the event annulled")]),
    ("Register Coverage Audit 2026 sample (register_coverage_audit_2026_sample.csv)",
     "One record per sampled person, as the register stood at the start of 1 January 2026.",
     [("AUDIT_NO", "audit case number"), ("PERSON_KEY", "person key"), ("MUN_CODE", "municipality of registration"),
      ("SEX", "as in the extract"), ("BIRTH_DATE", "as in the extract"), ("BIRTH_COUNTRY", "as in the extract"),
      ("NATIONALITY", "as in the extract"), ("ADDRESS_SINCE", "as in the extract")]),
    ("Register Coverage Audit 2026 visits (register_coverage_audit_2026_visits.csv)",
     "One record per fieldwork contact attempt or check; the last record of a case carries its final result.",
     [("AUDIT_NO", "audit case number"), ("VISIT", "attempt number within the case"),
      ("VISIT_DATE", "date (YYYY-MM-DD)"), ("MODE", "V visit, T telephone, R records check"),
      ("RESULT", "result code (list below)")]),
    ("Deregistration messages 2025 (deregistration_messages_2025.csv)",
     "The messages the in-play municipalities' register systems sent to the national hub for departures abroad, deaths "
     "and removals after an address investigation, as the hub logged them in 2025; the hub's log runs in UTC.",
     [("MESSAGE_ID", "message number: sending municipality, year, serial"), ("SENT_AT", "time sent (UTC)"),
      ("SENDER_MUN", "sending municipality, its code at sending"), ("PERSON_KEY", "person key"),
      ("MESSAGE_TYPE", "EMI departure abroad, OVL death, VOW removal after an address investigation"),
      ("EVENT_DATE", "date of departure, death or removal (YYYY-MM-DD)"),
      ("COUNTRY", "EMI: country of destination; OVL: country of death; VOW: 0000")]),
    ("Hub acknowledgements 2025 (hub_acknowledgements_2025.csv)",
     "The national hub's answer to each message. The Population Register records what an accepted message reports; "
     "a rejected message records nothing.",
     [("MESSAGE_ID", "the message answered"), ("RECEIVED_AT", "time received (UTC)"),
      ("RESULT", "ACCEPTED or REJECTED"), ("REASON", "rejection reason (list below)")]),
    ("Non-resident records (non_resident_records_extract.csv)",
     "Former residents of the in-play municipalities whose registration ended on departure abroad or on removal "
     "after an address investigation, as the municipality of last residence keeps them.",
     [("PERSON_KEY", "person key"), ("LAST_MUN_CODE", "municipality of last residence, its code at departure"),
      ("DEPARTURE_DATE", "date the registration ended (YYYY-MM-DD)"),
      ("COUNTRY", "country of destination; 0000 where the whereabouts are unknown"),
      ("STATUS", "A living abroad, O deceased"), ("STATUS_DATE", "date of the status (YYYY-MM-DD)"),
      ("SEX", "M or F"), ("BIRTH_DATE", "YYYYMMDD")]),
]
CODE_LISTS = [
    ("Event codes", [("GEB", "birth registered in the country"),
                     ("IMM", "settlement from abroad: the person is registered on arrival from abroad"),
                     ("EMI", "departure abroad: the registration ends on departure abroad"),
                     ("VER", "move from another municipality, registered by the new municipality"),
                     ("ADR", "change of address within the municipality"),
                     ("HUW", "marriage or registered partnership"), ("NAT", "acquisition of the nationality"),
                     ("OVL", "death: the registration ends on the date of death"),
                     ("ANN", "annulment of an earlier event: the annulled event did not take place and is no "
                             "part of the person's record")]),
    ("Audit result codes", [("10", "living at the registered address"),
                            ("11", "living elsewhere in the registered municipality"),
                            ("50", "living in the municipality on 1 January 2026, left afterwards"),
                            ("20", "living elsewhere in the country, not registered there"),
                            ("30", "living abroad"), ("40", "died before 1 January 2026, death not recorded"),
                            ("91", "no contact at this attempt (not a final result)")]),
    ("Hub rejection reasons", [(k, v) for k, v in sorted(__import__("sim").REASONS.items())]),
]


def write_layouts(path):
    h1, h1s, h2 = styles(NAVY)
    d, frame = _doc(path, "Population Register statistical extracts: record layouts", W.BUREAU,
                    "%s  |  Statistical extracts of the Population Register  |  record layouts, edition 2026" % W.BUREAU)
    s = [Paragraph("Statistical extracts of the Population Register", h1),
         Paragraph("Record layouts and code lists, edition 2026. %s, data services." % W.BUREAU, h1s)]
    for title, note, fields in LAYOUTS:
        s += [Paragraph(title, h2), Paragraph(note, SANS),
              _table([["Field", "Content"]] + [list(f) for f in fields], [4.2 * cm, 12.4 * cm]), Spacer(1, 4)]
    for title, codes in CODE_LISTS:
        s += [Paragraph(title, h2), _table([["Code", "Meaning"]] + [list(c) for c in codes], [2.2 * cm, 14.4 * cm]),
              Spacer(1, 4)]
    s.append(Paragraph("Country codes are those of the Bureau's country table (country_codes.csv); 0000 is unknown.",
                       SMALL))
    d.build(s, onFirstPage=frame, onLaterPages=frame)


# ------------------------------------------------------------------------------------ the delivery specification
SPEC_COLUMNS = [
    ("MUNICIPALITY_CODE", "GM followed by the four-digit municipality code"),
    ("MUNICIPALITY", "municipality name"),
    ("REGISTERED_20260101", "persons registered as residents of the municipality at the start of 1 January 2026"),
    ("NOT_RESIDENT_ESTIMATE", "of those persons, the number estimated from the Register Coverage Audit 2026 not to "
                              "have been living in the municipality on 1 January 2026, whole persons"),
    ("ADJUSTED_POPULATION", "REGISTERED_20260101 less NOT_RESIDENT_ESTIMATE: the adjusted resident population of "
                            "Article 14(2), whole persons"),
    ("QUALIFIES", "Y where ADJUSTED_POPULATION is below 5,000, otherwise N"),
    ("DEPARTURES_ABROAD_2025", "departures abroad in 2025 recorded in the Population Register from the messages the "
                               "municipality's register sent in 2025; a person who left twice counts twice, and a "
                               "removal after an address investigation is not a departure"),
    ("DEATHS_2025", "deaths in 2025 recorded in the Population Register from the messages the municipality's register "
                    "sent in 2025, of persons registered as residents of the municipality when they died"),
]


def write_spec(path):
    doc = docx.Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    t = doc.add_paragraph()
    r = t.add_run("Delivery specification: adjusted resident population and 2025 population account")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    doc.add_paragraph("%s, %s. Version 1.0, 1 October 2026. Contact: %s." % (W.MINISTRY, W.DIRECTORATE,
                                                                             W.PEOPLE["ministry"]))
    doc.add_heading("1. Purpose", level=2)
    doc.add_paragraph("The Minister decides in December 2026 which municipalities receive the small-municipality "
                      "supplement for 2027 under Article 14 of the Municipal Fund Regulation 2027. The ministry asks "
                      "the Bureau for the adjusted resident population of every municipality in play, with two "
                      "columns of the register's 2025 population account for the ministry's review of register "
                      "maintenance.")
    doc.add_heading("2. Municipalities in play", level=2)
    doc.add_paragraph("Every municipality whose register count at the start of 1 January 2026 was between 4,500 and "
                      "6,000 persons, as constituted on 1 January 2026: 24 municipalities. Ravenmeer includes its two "
                      "former municipalities, Ravensloot and Meerdijk.")
    doc.add_heading("3. File", level=2)
    doc.add_paragraph("adjusted_population_supplement_2027.csv: UTF-8, comma separated, one header row with the "
                      "column names below, one row per municipality in play, sorted by MUNICIPALITY_CODE. Counts are "
                      "whole numbers without separators. Dates and years in this specification are the country's "
                      "(Central European Time).")
    tab = doc.add_table(rows=1, cols=2)
    tab.style = "Table Grid"
    tab.rows[0].cells[0].text, tab.rows[0].cells[1].text = "Column", "Content"
    for name, text in SPEC_COLUMNS:
        row = tab.add_row().cells
        row[0].text, row[1].text = name, text
    doc.add_heading("4. Delivery", level=2)
    doc.add_paragraph("Deliver the file by 13 November 2026 through the ministry's secure exchange, with a chart of "
                      "the adjusted resident populations against the 5,000 line for the Minister's note.")
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading"):
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    doc.core_properties.title = "Delivery specification: adjusted resident population 2027"
    doc.core_properties.author = W.PEOPLE["ministry"]
    doc.save(path)


# ----------------------------------------------------------------------------------------------- the workbooks
def _head(ws, cols, widths):
    ws.append(cols)
    for i, w in enumerate(widths):
        c = ws.cell(row=1, column=i + 1)
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    ws.freeze_panes = "A2"


PROVINCE = {"0402": "Kustland", "0417": "Kustland", "0426": "Westmark", "0438": "Westmark",
            "0451": "Oosterland", "0465": "Westmark", "0473": "Oosterland", "0488": "Oosterland",
            "0496": "Middengouw", "0509": "Middengouw", "0514": "Kustland", "0527": "Middengouw",
            "0533": "Oosterland", "0546": "Middengouw", "0552": "Westmark", "1987": "Oosterland",
            "0568": "Kustland", "0571": "Westmark", "0584": "Middengouw", "0599": "Kustland",
            "0603": "Oosterland", "0615": "Westmark", "0628": "Middengouw", "0636": "Kustland",
            "0731": "Oosterland", "0744": "Oosterland"}


def write_codes(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Municipalities 2026"
    _head(ws, ["Code", "Municipality", "Province", "In play"], [9, 18, 16, 9])
    for code in sorted(W.CODES):
        ws.append([code, W.NAME[code], PROVINCE[code], "yes"])
    ws2 = wb.create_sheet("Changes 1 January 2026")
    _head(ws2, ["Former code", "Former municipality", "New code", "New municipality", "Effective"], [12, 20, 10, 18, 12])
    for name, code in W.FORMER:
        ws2.append([code, name, W.MERGED[1], W.MERGED[0], dt.date(2026, 1, 1)])
        ws2.cell(row=ws2.max_row, column=5).number_format = "yyyy-mm-dd"
    wb.save(path)


def write_pilot(path, seed=W.SEED):
    """the 2024 coverage pilot in two municipalities that are not in play (another reference date, another design)"""
    import numpy as np
    r = np.random.default_rng([seed, 120])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pilot results"
    _head(ws, ["Code", "Municipality", "Reference date", "Addresses visited", "Persons registered there",
               "Not living there", "Share"], [8, 16, 13, 12, 14, 12, 9])
    for code, name, addr in (("0212", "Grevenhoek", 610), ("0377", "Valdersum", 585)):
        persons = int(addr * float(r.uniform(2.05, 2.30)))
        gone = int(round(persons * float(r.uniform(0.012, 0.031))))
        ws.append([code, name, dt.date(2024, 1, 1), addr, persons, gone, gone / persons])
        ws.cell(row=ws.max_row, column=3).number_format = "yyyy-mm-dd"
        ws.cell(row=ws.max_row, column=7).number_format = "0.0%"
    ws.append([])
    ws.append(["The pilot visited a random sample of addresses in each municipality and checked every person "
               "registered at them on 1 January 2024. It tested the fieldwork procedures for the 2026 audit."])
    wb.save(path)
