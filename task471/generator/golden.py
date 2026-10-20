"""The golden deliverables, computed from the shipped pack alone: eli_2027_program_review.xlsx (the 2027 reviews, the
Form 26A conversion, the 2024-25 completers and the notes), eli_2026_program_pass_rates.png, and generator/figures.json
with every figure the write-up states."""
import json
import os
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt          # noqa: E402
import numpy as np                       # noqa: E402
import openpyxl                          # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import bundle as B   # noqa: E402
import world as W    # noqa: E402

GOLDEN = {"workbook": "eli_2027_program_review.xlsx", "chart": "eli_2026_program_pass_rates.png"}
W.GOLDEN = GOLDEN
PLUM = "4A2F5C"
FILL_H = PatternFill("solid", fgColor="E4DDEB")
FILL_R = PatternFill("solid", fgColor="F6E3E3")
THIN = Side(style="thin", color="A9A3B0")


def r1(a, n):
    """a rate in per cent to one decimal, half away from zero on the exact fraction"""
    return float(np.floor(1000.0 * a / n + 0.5) / 10.0)


def compute(target):
    pk = B.Pack(target)
    g, e = pk.cut("golden")
    fam = pk.cut("golden", "all")
    comp = pk.completers()
    rates = pk.rates(g, comp)
    prior = pk.prior_year()
    progs = []
    for k, pr in enumerate(W.PROGRAMS):
        n, p = rates[k]
        n0, p0 = prior[k]
        progs.append(dict(code=pr[1], name=pr[0], counted=n, passed=p, rate=r1(p, n), rate_exact=p / n,
                          review=(p / n < W.REVIEW_LINE), prior_tested=n0, prior_passed=p0, prior_rate=r1(p0, n0)))
    review = [q["name"] for q in progs if q["review"]]
    above = min((q for q in progs if not q["review"]), key=lambda q: q["rate_exact"])
    below = max((q for q in progs if q["review"]), key=lambda q: q["rate_exact"])
    X, VX, VXm, Y, VY, VYm = pk.arrays()
    bias = pk.validity_bias()
    conv = []
    for raw in range(W.MAX_RAW + 1):
        eq = float(e(raw))
        rr = int(np.clip(np.floor(eq + 0.5), 0, W.MAX_RAW))
        conv.append([raw, round(eq, 2), rr, pk.conv[rr], "Pass" if pk.conv[rr] >= W.PASS_SCALE else None])
    rungs = {}
    for route in ("as scored", "multiple-choice anchor", "every assignment's papers"):
        rungs[route] = sorted({v[0] for v in pk.cut(route, "all").values()})
    rungs = {k: (v[0] if len(v) == 1 else v) for k, v in rungs.items()}
    rungs["no equating"] = pk.cut("no equating")[0]
    rungs["raters adjusted, no equating"] = pk.cut("raters adjusted, no equating")[0]
    as_scored = pk.rates(rungs["as scored"], comp)
    rung_sets = {"as scored": [W.PROGRAMS[k][0] for k, (n, p) in as_scored.items() if p / n < W.REVIEW_LINE]}
    out = dict(
        review=review, review_count=len(review), raw_cut=g, programs=progs, rungs=rungs, rung_sets=rung_sets,
        nearest_above=dict(name=above["name"], rate=above["rate"]), nearest_below=dict(name=below["name"], rate=below["rate"]),
        validity_bias=round(bias, 4), anchor_shift=round(-2 * bias * W.WA_WEIGHT, 2),
        validity_bias_below_2025=round(-bias, 4), wa2_score_shift=round(-2 * bias, 4), rule="Rule 8.4",
        wa2_papers=sum(1 for p in pk.papers.values() if p["wa"] == "WA2"),
        wa2_validity_reads=sum(1 for p, a, v in pk.vreads if a == "WA2"),
        n_first_2025=int(len(Y)), n_first_2026=int(len(X)),
        equated_at_cut={k: round(v[1], 2) for k, v in fam.items()},
        equated_below_cut={k: round(v[2], 2) for k, v in fam.items()},
        counted_total=sum(q["counted"] for q in progs), passed_total=sum(q["passed"] for q in progs),
        prior_tested_total=sum(q["prior_tested"] for q in progs), prior_passed_total=sum(q["prior_passed"] for q in progs),
        conversion=conv, pass_raw_25a=W.PASS_RAW_25A, pass_scale=W.PASS_SCALE,
        call=sorted(review),
        ask={q["code"]: dict(tested=q["prior_tested"], rate=q["prior_rate"]) for q in progs})
    return out


def write_workbook(path, fig):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reviews 2027"
    words = {3: "Three", 4: "Four", 5: "Five", 6: "Six", 7: "Seven"}
    ws["A1"] = "%s of sixteen elementary programs fall below 70 per cent and go to formal review in 2027" % (
        words.get(fig["review_count"], str(fig["review_count"])))
    ws["A1"].font = Font(bold=True, size=13, color=PLUM)
    ws["A2"] = ("First-attempt pass rates of 2025-26 completers, ELI 104, Rule 8.4. Office of Assessment Research, "
                "August 21, 2026. Form 26A raw passing score: %d (equated to Form 25A's raw %d, scale 240)." % (
                    fig["raw_cut"], fig["pass_raw_25a"]))
    ws["A2"].font = Font(italic=True, size=9, color="5B6067")
    head = ["Program code", "Program", "Completers counted", "Passed on first attempt", "First-attempt pass rate (%)",
            "Formal review in 2027"]
    ws.append([])
    ws.append(head)
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = FILL_H
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = Border(bottom=THIN)
    for q in fig["programs"]:
        ws.append([q["code"], q["name"], q["counted"], q["passed"], q["rate"], "Yes" if q["review"] else "No"])
        if q["review"]:
            for c in ws[ws.max_row]:
                c.fill = FILL_R
    ws.append([])
    ws.append([None, "All sixteen programs", fig["counted_total"], fig["passed_total"],
               float(np.floor(1000.0 * fig["passed_total"] / fig["counted_total"] + 0.5) / 10.0), None])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=5, min_col=5, max_col=5):
        for c in row:
            c.number_format = "0.0"
    for col, wdt in zip("ABCDEF", (13, 32, 12, 13, 14, 12)):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "C5"
    cv = wb.create_sheet("Form 26A conversion")
    cv.append(["Form 26A raw score", "Equated Form 25A raw score", "Rounded", "Scale score", "Result"])
    for c in cv[1]:
        c.font = Font(bold=True)
        c.fill = FILL_H
        c.alignment = Alignment(wrap_text=True)
    for row in fig["conversion"]:
        cv.append(row)
    for rr in cv.iter_rows(min_row=2, min_col=2, max_col=2):
        for c in rr:
            c.number_format = "0.00"
    for col, wdt in zip("ABCDE", (12, 14, 10, 10, 9)):
        cv.column_dimensions[col].width = wdt
    cv.freeze_panes = "A2"
    pr = wb.create_sheet("2024-25 completers")
    pr.append(["Program code", "Program", "2024-25 completers with a first attempt by June 14, 2025",
               "Passed on first attempt", "First-attempt pass rate (%)"])
    for c in pr[1]:
        c.font = Font(bold=True)
        c.fill = FILL_H
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for q in fig["programs"]:
        pr.append([q["code"], q["name"], q["prior_tested"], q["prior_passed"], q["prior_rate"]])
    for rr in pr.iter_rows(min_row=2, min_col=5, max_col=5):
        for c in rr:
            c.number_format = "0.0"
    for col, wdt in zip("ABCDE", (13, 32, 16, 13, 14)):
        pr.column_dimensions[col].width = wdt
    pr.freeze_panes = "C2"
    nt = wb.create_sheet("Notes")
    lines = [
        ("How these figures were made", True),
        ("Completers: each program's Title II list for 2025-26, elementary pathway only (Elementary, or Elementary and "
         "Special Education), completion dates July 1, 2025 to June 30, 2026, each educator once (Rule 8.4.2).", False),
        ("First attempt: the educator's earliest attempt in the licensing system's history, voided attempts excluded, "
         "at or before June 13, 2026 (8.4.3, 8.4.4). A June 2025 or earlier first attempt keeps its reported result.", False),
        ("June 2026 attempts are placed on the reporting scale through Form 25A: Form 26A is equated to Form 25A under the "
         "common-item nonequivalent groups design through its 20 shared items and WA2, on each form's first-attempt "
         "candidates (%s for 25A, %s for 26A)." % ("{:,}".format(fig["n_first_2025"]), "{:,}".format(fig["n_first_2026"])),
         False),
        ("WA2 is read by a different scoring team each year. Arcwright's raters read the %d WA2 validity papers, June "
         "2025 responses, %.2f points per read below the scores their two June 2025 reads agreed on (%d reads), so each "
         "2026 WA2 score is moved up by twice that before it enters the anchor (%.2f anchor points, the assignment "
         "counting twice)." % (
             fig["wa2_papers"], -fig["validity_bias"], fig["wa2_validity_reads"], fig["anchor_shift"]), False),
        ("Equated scores are rounded to whole numbers and converted with Form 25A's table; a scale score of 240 passes "
         "(8.4.5). Tucker, Levine and chained linear equating at each synthetic weight and chained equipercentile "
         "equating all put the Form 26A passing raw score at %d." % fig["raw_cut"], False),
        ("2024-25 completers: the Board's completer register, elementary pathway, program year 2024-25, with a first "
         "attempt at or before June 14, 2025. Where a completer's register record has no attempts, the attempt is found "
         "under the educator record with the same name and birth date.", False),
        ("Rates are rounded to one decimal; a rate below 70.0 per cent on the exact fraction goes to review.", False)]
    for i, (t, b) in enumerate(lines, 1):
        c = nt.cell(row=i, column=1, value=t)
        c.font = Font(bold=b, size=12 if b else 10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    nt.column_dimensions["A"].width = 115
    wb.save(path)


def write_chart(path, fig):
    progs = sorted(fig["programs"], key=lambda q: q["rate_exact"])
    names = [q["name"] for q in progs]
    vals = [q["rate"] for q in progs]
    colors = ["#9b2d3c" if q["review"] else "#7d8a96" for q in progs]
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9})
    fig_, ax = plt.subplots(figsize=(8.6, 6.4), dpi=120)
    y = np.arange(len(progs))
    ax.barh(y, vals, color=colors, height=0.66)
    ax.axvline(70, color="#1b1d1f", lw=1.2, ls="--")
    ax.text(70.4, len(progs) - 0.45, "70 per cent: formal review below the line", fontsize=8.5, va="bottom")
    for i, (v, q) in enumerate(zip(vals, progs)):
        if q["review"]:
            ax.text(v - 0.5, i, "%.1f" % v, va="center", ha="right", fontsize=8, color="white", fontweight="bold")
        else:
            ax.text(v + 0.4, i, "%.1f" % v, va="center", fontsize=8)
    ax.set_yticks(y)
    ax.set_yticklabels(names)
    ax.set_xlim(40, 100)
    ax.set_xlabel("First-attempt pass rate of 2025-26 completers, ELI 104 (per cent)")
    words = {3: "Three", 4: "Four", 5: "Five", 6: "Six", 7: "Seven"}
    ax.set_title("%s programs fall below 70 per cent on the June 2026 test" % words.get(fig["review_count"],
                                                                                      str(fig["review_count"])), loc="left",
                 fontsize=11, fontweight="bold", color="#4A2F5C")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.text(0, -0.09, "Form 26A equated to Form 25A (raw passing score %d). Office of Assessment Research, August 21, "
            "2026." % fig["raw_cut"], transform=ax.transAxes, fontsize=7.5, color="#5b6067")
    fig_.tight_layout()
    fig_.savefig(path, metadata={"Software": None})
    plt.close(fig_)


def main(target, gdir, fig_path):
    os.makedirs(gdir, exist_ok=True)
    fig = compute(target)
    write_workbook(os.path.join(gdir, GOLDEN["workbook"]), fig)
    write_chart(os.path.join(gdir, GOLDEN["chart"]), fig)
    import postprocess as PP
    PP.normalize_ooxml(os.path.join(gdir, GOLDEN["workbook"]), "Office of Assessment Research",
                       "Office of Assessment Research", "2026-08-20T14:02:00Z", "2026-08-21T16:47:31Z", excel=True,
                       company=W.BOARD)
    with open(fig_path, "w", encoding="utf-8") as fh:
        json.dump(fig, fh, indent=1)
        fh.write("\n")
    return fig
