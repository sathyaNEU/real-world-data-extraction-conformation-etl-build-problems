"""The shipped pack read back into the estimators' data structure (methods.py), from the files alone: the census
workbook, the survey's two person files, the roll on census day and its electorates, the vital and border series and
the roll's transaction log with the electorates of the 2023 general election's roll."""
import csv
import os

import openpyxl

import methods as M
import pack as P
import world as W


def rows(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def _period(p):
    return "2025Q" if p == P.PERIOD_2025 else int(p)


class Pack:
    def __init__(self, target):
        self.target = target
        f = lambda k: os.path.join(target, P.FILES[k])
        conc = rows(f("electorates"))
        self.region_of = {r["ELECTORATE"]: r["REGION"] for r in conc}
        self.region_2020 = {r["ELECTORATE"]: r["REGION"] for r in rows(f("roll_2023"))}
        d = {"census": {}, "e": {}, "p": {}, "roll": {c: 0 for c in W.CODES}, "births": {}, "deaths": {},
             "departures": {}, "arrivals": {}}
        wb = openpyxl.load_workbook(f("census"), read_only=True)
        for c in W.CODES:
            ws = wb["%s %s" % (c, W.NAME[c])]
            for r in ws.iter_rows(min_row=5, max_row=105, values_only=True):
                age = 100 if r[0] == "100 and over" else int(r[0])
                d["census"][(c, "M", age)] = (int(r[1]), int(r[2]))
                d["census"][(c, "F", age)] = (int(r[3]), int(r[4]))
        for r in rows(f("pes_records")):
            key = (r["REGION"], r["SEX"], min(int(r["AGE"]), 100))
            e = d["e"].setdefault(key, [0, 0])
            e[0] += 1
            e[1] += int(r["STATUS"] == "C")
        for r in rows(f("pes_people")):
            key = (r["REGION"], r["SEX"], min(int(r["AGE"]), 100))
            p = d["p"].setdefault(key, [0, 0])
            p[0] += 1
            p[1] += int(r["MATCH"] == "M")
        self.snapshot = rows(f("roll"))
        for r in self.snapshot:
            if r["ADDRESS_TYPE"] == "R":
                d["roll"][self.region_of[r["ELECTORATE"]]] += int(r["ELECTORS"])
        for r in rows(f("births")):
            d["births"][(_period(r["PERIOD"]), r["SEX"])] = int(r["LIVE_BIRTHS"])
        hist = openpyxl.load_workbook(f("deaths_hist"), read_only=True)
        for sex, title in (("M", "Males"), ("F", "Females")):
            ws = hist[title]
            head = None
            for r in ws.iter_rows(min_row=4, values_only=True):
                if head is None:
                    head = list(r)
                    continue
                if r[0] is None:
                    continue
                for c, v in zip(head[1:], r[1:]):
                    if v:
                        d["deaths"][(int(r[0]), sex, int(c))] = int(v)
        for r in rows(f("deaths")):
            d["deaths"][(_period(r["PERIOD"]), r["SEX"], int(r["YEAR_OF_BIRTH"]))] = int(r["DEATHS"])
        for r in rows(f("movements")):
            k = (_period(r["PERIOD"]), r["SEX"], int(r["YEAR_OF_BIRTH"]))
            if int(r["ARRIVALS"]):
                d["arrivals"][k] = int(r["ARRIVALS"])
            if int(r["DEPARTURES"]):
                d["departures"][k] = int(r["DEPARTURES"])
        self.d = d
        self.log = rows(f("log"))
        self.codes = list(W.CODES)

    # ------------------------------------------------------------------ the call
    def estimates(self, **kw):
        return M.dse_cells(self.d, **kw)

    def adults(self, rung="golden", **kw):
        est = self.estimates(**kw)
        if rung == "golden":
            return M.adults(est, M.factors(self.d, est))
        return M.adults(est)

    def gap(self, adults):
        return {c: adults[c] - self.d["roll"][c] for c in W.CODES}

    def allocation(self, adults=None):
        adults = adults or self.adults()
        return M.allocate(self.gap(adults))

    # ------------------------------------------------------------------ the asks
    def region(self, vintage=True):
        """a transaction's region: the old platform's electorate numbers are the 2020 review's, as on the 2023 general
        election's roll (vintage), or every number read as the 2024 review's"""
        def f(r):
            if vintage and M.old_platform(r):
                return self.region_2020[r["ELECTORATE"]]
            return self.region_of[r["ELECTORATE"]]
        return f

    def new_electors(self, vintage=True, **kw):
        return M.new_electors(self.log, self.region(vintage), **kw)

    def removed_dead(self, vintage=True, **kw):
        return M.removed_dead(self.log, self.region(vintage), **kw)
