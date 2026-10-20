"""The generator's assertions on the built pack: input gates, the frames and linkage against the simulation, the call and
its convergence (C1) across the design-based estimators, every rung of the ladder, the three-rules block (the analyst
test on the clean truth), every row of the semantic fork table, the ask's stops and the separation line, and
figures.json against the golden files."""
import json
import os
import re

import numpy as np
import openpyxl

import bundle as B
import pack as P
import sim as S
import world as W

GOLD_SET = ["C01", "C02", "C06", "C07", "C11", "C13"]
RUNG_SETS = {"pooled": ["C01", "C07", "C13"],
             "drawn list": ["C01", "C02", "C04", "C06", "C07", "C09", "C10", "C11", "C13", "C16"],
             "lists added": ["C01", "C02", "C04", "C06", "C07", "C09", "C10", "C11", "C13", "C16"],
             "permit list": ["C01", "C07", "C13"],
             "parcel list": ["C02", "C06", "C07", "C11", "C13"]}


class Checker:
    def __init__(self):
        self.ok = True
        self.n = 0

    def __call__(self, cond, msg):
        self.n += 1
        if not cond:
            self.ok = False
            print("FAIL:", msg)


def _pdf_text(path):
    import subprocess
    return subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True).stdout


def run(target, gdir, x=None, meta_path=None, fig_path=None):
    ck = Checker()
    x = x or P.everything()
    pk = B.Pack(target)
    sim = {c["code"]: c for c in x["counties"]}
    # ------------------------------------------------------------------------------------------- input gates
    files = sorted(os.listdir(target))
    ck(len(files) >= 10, "ten or more files")
    ck(len({f.rsplit(".", 1)[1] for f in files}) >= 3, "three or more formats")
    with open(os.path.join(target, P.FILES["permits"]), encoding="utf-8") as fh:
        ck(sum(1 for _ in fh) - 1 >= 25000, "a file of 25,000 or more rows")
    if meta_path:
        meta = json.load(open(meta_path, encoding="utf-8"))
        ck(len(meta["distractor_files"]) >= 2 and all(f in files for f in meta["distractor_files"]),
           "two or more distractors named in metadata.json")
    # ------------------------------------------------------------------------- frames, linkage, the sample
    for k in pk.codes:
        c, L, s = sim[k], pk.link[k], pk.sample[k]
        ck(L["NP"] == c["NP"] and L["NQ"] == c["NQ"], "%s frame sizes" % k)
        ck(L["NB"] == int((c["grp"] == 2).sum()) and L["union"] == c["N"], "%s linkage: wells on both, union" % k)
        ck(s["nP"] == c["nP"] and s["nQ"] == c["nQ"], "%s targets met from the manifest" % k)
        truth = {c["kit"][i]: i for i in c["kit"]}
        bad = 0
        for kit, w in s["wells"].items():
            i = truth[kit]
            bad += (w["onP"] != bool(c["onP"][i])) + (w["onQ"] != bool(c["onQ"][i]))
            bad += (w["y"] != bool(c["measured"][i] > W.STANDARD_N))
        ck(bad == 0, "%s each tested well's lists and first test as simulated" % k)
    # ------------------------------------------------------------------------------- the call and C1
    gold = pk.qualifying()
    ck(gold == GOLD_SET, "the six counties: %s" % gold)
    for r, kw in B.GOLD_ROUTES:
        ck(pk.qualifying(r, **kw) == GOLD_SET, "C1: %s %s gives the same six" % (r, kw))
    for k in pk.codes:
        vals = [pk.estimate(k, r, **kw) for r, kw in B.GOLD_ROUTES]
        if k in GOLD_SET:
            ck(min(vals) >= 1.15 * W.QUALIFY, "%s qualifies by 15 per cent on every estimator (%.0f)" % (k, min(vals)))
        else:
            ck(max(vals) <= 0.90 * W.QUALIFY, "%s short by 10 per cent on every estimator (%.0f)" % (k, max(vals)))
    # ------------------------------------------------------------------------------------------ the ladder
    for r, want in RUNG_SETS.items():
        got = pk.qualifying(r)
        ck(got == want, "rung %s: %s" % (r, got))
        ck(got != GOLD_SET, "rung %s names its own set" % r)
    # ---------------------------------------------- the three rules: the analyst test on the clean truth
    clean = {k: S.estimates(sim[k]) for k in pk.codes}
    easy = [k for k in pk.codes if clean[k]["pooled"] >= W.QUALIFY]
    right = [k for k in pk.codes if clean[k]["KA"] >= W.QUALIFY]
    ck(easy == RUNG_SETS["pooled"], "analyst test: the pooled share on the clean truth files the three (%s)" % easy)
    ck(right == GOLD_SET, "analyst test: the decisive step alone on the clean truth gives the six (%s)" % right)
    ck([k for k in pk.codes if clean[k]["drawn list's rate"] >= W.QUALIFY] == RUNG_SETS["drawn list"],
       "analyst test: the drawn list's rate on the clean truth files the ten")
    # ------------------------------------------------------------------------------- semantic fork table
    design = _pdf_text(os.path.join(target, P.FILES["design"]))
    standard = _pdf_text(os.path.join(target, P.FILES["standard"]))
    flat_d, flat_s = " ".join(design.split()), " ".join(standard.split())
    # SF1 "above the standard": more than 10 against 10 or more; no result sits on 10
    on = [res for res in pk.results.values() if abs(res["n"] - W.STANDARD_N) < 1e-9]
    ck(not on, "SF1: no result is exactly 10 mg/L, so above and at-or-above select the same wells")
    ck("more than 10 mg/L nitrate as nitrogen" in flat_s, "SF1: the standard defines above")
    # SF2 the survey population as it stood on August 31, 2025, against the extract's own date
    alt = B.Pack(target, frame_date="2026-12-31")
    ck(alt.qualifying() == GOLD_SET, "SF2: domestic wells permitted after the frame date included: same call")
    ck("as they stood on August 31, 2025" in flat_d, "SF2: the design report files the frame date")
    # SF3 a reanalysed sample: the later report against the earlier one
    alt = B.Pack(target, results="earliest")
    ck(alt.qualifying() == GOLD_SET, "SF3: the earlier reports of reanalysed samples: same call")
    ck("its later report replaces the earlier one" in flat_d, "SF3: the design report files the later report")
    # SF4 Kinsale's NO3: the design report gives the factor; read as nitrogen the call moves
    ck("44.27 mg/L nitrate as NO3" in flat_d, "SF4: the design report converts Kinsale's NO3")
    # SF12 Kinsale's factor: the design report's 4.427 against the fact sheet's 4.4 (and its 44 mg/L): no Kinsale
    # nitrate sits between 44.0 and 44.3 mg/L NO3, so both put every result on the same side
    import csv as _csv
    with open(os.path.join(target, P.FILES["kin"]), encoding="utf-8", newline="") as fh:
        kin = [float(r["VALUE"]) for r in _csv.DictReader(fh) if r["PARAMETER"] == "NITRATE" and r["FLAG"] != "<"]
    ck(not [v for v in kin if 43.9 <= v <= 44.4], "SF12: no Kinsale result between 43.9 and 44.4 mg/L NO3")
    # SF5 250 or more against more than 250
    ests = [pk.estimate(k, r, **kw) for k in pk.codes for r, kw in B.GOLD_ROUTES]
    ck(min(abs(e - W.QUALIFY) for e in ests) > 1, "SF5: no estimate within a well of 250")
    # SF6 the linkage: parcel number or address; by address alone the same wells link (C1), by parcel number alone a
    # permit without one links nothing, which the filed sentence rules out
    alt = B.Pack(target, link="address only")
    ck(alt.qualifying() == GOLD_SET and all(alt.link[k]["NB"] == pk.link[k]["NB"] for k in pk.codes),
       "SF6: linkage by address alone: the same wells on both lists, same call")
    alt = B.Pack(target, link="parcel only")
    ck(alt.qualifying() != GOLD_SET, "SF6: linkage by parcel number alone moves the call (%s), so the sentence "
       "carries it" % alt.qualifying())
    ck("when the permit's parcel number is the parcel's or, on a permit that carries no parcel number, when the "
       "permit's site address is the parcel's situs address" in flat_d, "SF6: the design report files the linkage")
    # SF7 the estimate covers the whole survey population, not the tested wells counted
    ck(max(pk.counts(k)["above"] for k in pk.codes) < W.QUALIFY, "SF7: no county has 250 tested wells above")
    ck("the wells the Survey did not test as well as those it did" in flat_s, "SF7: the standard files the population")
    # SF8 a well tested earlier counts toward the parcel list's target as well
    alt = B.Pack(target, earlier="permit list only")
    ck(alt.qualifying() == GOLD_SET, "SF8: TESTED_EARLIER left off the parcel list's sample: same call (%s)" %
       alt.qualifying())
    ck("counts toward the parcel list's target as well" in flat_d, "SF8: the design report files it")
    # SF9 a blank water code is not a private well
    ck(all(p["code"] not in B.PRIVATE for k in pk.codes for p in pk.parcel_all[k] if p["code"] == ""),
       "SF9: blanks are off the parcel list")
    # SF10 the estimate's rounding: the verdict is on the unrounded estimate and no estimate sits within half a well
    ck(min(abs(e - W.QUALIFY) for e in ests) > 0.5, "SF10: rounding the estimate cannot move a county")
    # SF11 county: a well's county is the county of its permit or parcel; the manifest's county agrees
    ck(all(r["COUNTY_CODE"] == k for k in pk.codes for r in pk.draws[k]), "SF11: manifest counties")
    # ------------------------------------------------------------------------------------------- the ask
    gold_ret = pk.retests()
    truth = {}
    close = np.datetime64(W.RETESTS_CLOSE.isoformat())
    for c in x["counties"]:
        qs = [q for q in c["requests"] if q["kind"] == "above" and q["collected"] <= close]
        truth[c["code"]] = [len(qs), sum(1 for q in qs if q["value"] > W.STANDARD_N)]
    ck(gold_ret == truth, "ask: the golden handling reproduces the truth exactly")
    stops = {"by request": pk.retests(by="request"), "portal only": pk.retests(channel="portal"),
             "no window": pk.retests(window=False), "no first-above": pk.retests(first_above=False)}
    moved = stops["by request"]
    ck(all(moved[k][0] != gold_ret[k][0] and moved[k][1] != gold_ret[k][1] for k in pk.codes),
       "ask primary: the request's county moves all 32 figures")
    for name, st in stops.items():
        hits = sum(st[k][j] == gold_ret[k][j] for k in pk.codes for j in (0, 1))
        print("ask stop %-15s keeps %2d of 32" % (name, hits))
    ck(sum(stops["portal only"][k][0] == gold_ret[k][0] for k in pk.codes) == 0, "ask hazard: phone requests move every county")
    # separation: the call reads no request and no retest result
    kits = {r["RETEST_KIT"] for r in pk.requests}
    first_kits = {w for k in pk.codes for w in pk.sample[k]["wells"]}
    ck(not (kits & first_kits), "separation: no retest kit is a first-test kit")
    # ------------------------------------------------------------------------------------ figures and golden
    if fig_path:
        fig = json.load(open(fig_path, encoding="utf-8"))
        ck(fig["call_codes"] == GOLD_SET and fig["qualifying_count"] == 6, "figures.json call")
        wb = openpyxl.load_workbook(os.path.join(gdir, "nitrate_assistance_counties_2027.xlsx"), read_only=True)
        rows = [r for r in wb["Counties 2027"].iter_rows(min_row=5, values_only=True) if r[0]]
        got = {r[0]: r for r in rows}
        for r in fig["counties"]:
            g = got[r["code"]]
            ck(g[2] == r["union"] and (g[3] == "Yes") == r["qualifies"] and g[4] == r["retested"]
               and g[5] == r["still_above"], "golden workbook row %s" % r["code"])
            ck(r["union"] == pk.link[r["code"]]["union"], "figures union %s" % r["code"])
    # ---------------------------------------------------------------------------------------- house rules
    for f in files:
        if f.endswith(".pdf"):
            ck(not re.search("[\u2013\u2014]", _pdf_text(os.path.join(target, f))), "no dashes in %s" % f)
    # H25: every PDF reads the same through a strict parser as through pdftotext
    import csv as csvm
    import pypdf
    for f in files:
        if f.endswith(".pdf"):
            strict = " ".join(" ".join(p.extract_text() for p in pypdf.PdfReader(os.path.join(target, f),
                                                                                   strict=True).pages).split())
            loose = " ".join(_pdf_text(os.path.join(target, f)).split())
            words = re.findall(r"[A-Za-z]{6,}", loose)[:400]
            ck(sum(1 for w in words if w in strict) >= 0.97 * len(words), "H25: %s reads under a strict parser" % f)
            if f == P.FILES["design"]:
                ck("as they stood on August 31, 2025" in strict, "H25: the frame date reads strictly")
    # H16: no file's records run past the date the file carries
    import build as BD

    def latest(key, cols):
        with open(os.path.join(target, P.FILES[key]), encoding="utf-8") as fh:
            v = max(max(r[c] for c in cols) for r in csvm.DictReader(fh))
        return "%s-%s-%s" % (v[:4], v[4:6], v[6:8]) if v.isdigit() else v
    wbc = openpyxl.load_workbook(os.path.join(target, P.FILES["ced"]), read_only=True)
    ced = max(r[9] for r in wbc["Results"].iter_rows(min_row=2, values_only=True)).strftime("%Y-%m-%d")
    for k, v in (("hal", latest("hal", ["COLLECTED", "RECEIVED", "REPORTED"])),
                 ("kin", latest("kin", ["SAMPLED", "REPORTED"])), ("ced", ced),
                 ("retests", latest("retests", ["REQUESTED", "KIT_MAILED"])),
                 ("permits", latest("permits", ["DATE_DRILLED"])), ("manifest", latest("manifest", ["MAILED"]))):
        ck(v <= BD.AS_OF[k], "H16: %s runs to %s, on or before its date %s" % (k, v, BD.AS_OF[k]))
    print("checks: %d assertions, %s" % (ck.n, "all pass" if ck.ok else "FAILURES"))
    return ck.ok
