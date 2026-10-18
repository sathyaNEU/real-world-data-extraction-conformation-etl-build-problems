"""task248 goldens: the coordinator's workbook (each plant's winter 2027 episodes on the published basis and on the shed
load, the decay fit, the historian's hours without a valid inlet flow reading, the faecal-strength check's failures) and
the planning note to the surveillance manager with its chart."""
import io
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import params as P
import postprocess

XLSX = "winter_2027_episodes_shed_load.xlsx"
DOCX = "verification_desk_winter_2028.docx"
ORG = P.SERVICE
AUTHOR = P.COORDINATOR[0]


def _head(ws, row, head, height=46):
    thin = Side(style="thin", color="A6ADB8")
    for j, h in enumerate(head, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(bold=True, size=9.5)
        c.fill = PatternFill("solid", fgColor="DDE4EE")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = Border(bottom=thin)
    ws.row_dimensions[row].height = height


def _note(ws, row, text, ncol, height):
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(size=8.5, italic=True)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
    ws.row_dimensions[row].height = height


def _widths(ws, widths):
    for k, w in enumerate(widths):
        ws.column_dimensions[chr(65 + k)].width = w


def _workbook(path, G):
    F = G["fig"]
    plants = F["plants"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Episodes"
    ws["A1"] = "Winter 2027 alert episodes on the published basis and on the shed load"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Episodes opening in the weeks starting 3 May to 27 September 2027, by the dashboard's rules (methods note "
                "v4, sections 4 and 5). Plants in the register's order.")
    ws["A2"].font = Font(italic=True, size=9)
    head = ["Plant", "Name", "Travel time (min)", "Episodes, published", "Episodes, shed load", "Change"]
    _head(ws, 4, head, 44)
    r = 5
    for p in plants:
        vals = [p["code"], p["name"], p["tt"], p["published"], p["shed"], p["change"] or None]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            c.alignment = Alignment(horizontal="left" if j <= 2 else "right")
            if j == 6 and v:
                c.number_format = "+0;-0"
            if p["change"]:
                c.fill = PatternFill("solid", fgColor="FBEFD9")
        r += 1
    for j, v in enumerate(["Total", "", None, F["totals"]["published"], F["totals"]["shed"],
                           F["totals"]["shed"] - F["totals"]["published"]], 1):
        c = ws.cell(row=r, column=j, value=v)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="left" if j <= 2 else "right")
        if j == 6:
            c.number_format = "+0;-0"
    _note(ws, r + 2, "The number to plan for: %d alert episodes, last winter's samples read on the shed load. Six plants "
                     "move: %s." % (F["call"], ", ".join("%s %+d" % (c, v) for c, v in F["moves"].items())), 6, 30)
    _note(ws, r + 3, "Shed load: each sample's N1 load per person with the RNA lost in the sewer restored, exp(k(T) x "
                     "travel time), k(T) the paired study's decay rate at the composite's mean inlet temperature (sheet "
                     "Decay); weekly values, baselines and levels recomputed on it (section 6.1).", 6, 42)
    _widths(ws, [8, 22, 11, 12, 12, 9])
    ws.freeze_panes = "C5"

    dc = wb.create_sheet("Decay")
    dc["A1"] = "Decay of SARS-CoV-2 RNA in the trunk sewers, from the paired upstream study (2026)"
    dc["A1"].font = Font(bold=True, size=11)
    d = F["decay"]
    rows = [("Paired visits", d["visits"], "0"), ("k at 20 degrees C (per hour)", float(d["k20"]), "0.000"),
            ("theta (per degree C)", float(d["theta"]), "0.00")]
    for i, (a, b, fmt) in enumerate(rows, 3):
        dc.cell(row=i, column=1, value=a)
        c = dc.cell(row=i, column=2, value=b)
        c.number_format = fmt
    n = 3 + len(rows) + 1
    _head(dc, n, ["Study plant", "Travel time (min)", "Loss, December to February (log10)",
                  "Loss, June to August (log10)"], 44)
    tt = {p["code"]: p["tt"] for p in plants}
    nm = {p["code"]: p["name"] for p in plants}
    for i, code in enumerate(("DNK", "NTL", "ADS", "CLF"), n + 1):
        s = F["seasonal"][code]
        for j, v in enumerate([nm[code], tt[code], s["December to February"], s["June to August"]], 1):
            c = dc.cell(row=i, column=j, value=v)
            if j >= 3:
                c.number_format = "0.00"
    _note(dc, n + 6, "Fit: least squares of ln(plant N1 / upstream N1) on -k20 x theta^(T - 20) x travel time (hours), "
                     "the three replicates of each point averaged, T the mean of the two points' water temperatures. "
                     "Loss: the mean of log10(upstream / plant) over the season's visits.", 4, 42)
    _widths(dc, [30, 14, 18, 18])

    hs = wb.create_sheet("Historian hours")
    hs["A1"] = "Hours from 1 April 2027 to 31 March 2028 with no valid inlet flow reading in the historian"
    hs["A1"].font = Font(bold=True, size=11)
    _head(hs, 3, ["Plant", "Name", "No row", "Blank flow", "Held reading", "Zero or below", "Hours without a valid "
                                                                                          "reading"], 44)
    for i, p in enumerate(plants, 4):
        x = p["hours_detail"]
        for j, v in enumerate([p["code"], p["name"], x["no_row"], x["blank"], x["held"], x["zero_or_below"],
                               p["hours"]], 1):
            c = hs.cell(row=i, column=j, value=v)
            c.alignment = Alignment(horizontal="left" if j <= 2 else "right")
            c.font = Font(bold=(j == 7))
    m = 4 + len(plants)
    tot = [sum(p["hours_detail"][k] for p in plants) for k in ("no_row", "blank", "held", "zero_or_below")]
    for j, v in enumerate(["Total", ""] + tot + [F["ask_a_total"]], 1):
        c = hs.cell(row=m, column=j, value=v)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="left" if j <= 2 else "right")
        if j >= 3:
            c.number_format = "#,##0"
    m += 1
    hs.freeze_panes = "C4"
    _note(hs, m + 1, "The year's %s real local hours (the clock repeats 02:00 on 4 April 2027 and skips 02:00 on 26 "
                     "September 2027), exact duplicate rows of the export removed. Held: every reading after the first of "
                     "four or more identical consecutive readings (the export note)." % "{:,}".format(
                         F["conventions"]["real_hours"]), 7, 42)
    _widths(hs, [8, 22, 9, 9, 10, 10, 14])

    fs = wb.create_sheet("Faecal-strength check")
    fs["A1"] = "Composites from 4 October 2027 to 30 March 2028 that fail the faecal-strength check"
    fs["A1"].font = Font(bold=True, size=11)
    _head(fs, 3, ["Plant", "Name", "Composites failing"], 30)
    for i, p in enumerate(plants, 4):
        for j, v in enumerate([p["code"], p["name"], p["flags"]], 1):
            c = fs.cell(row=i, column=j, value=v)
            c.alignment = Alignment(horizontal="left" if j <= 2 else "right")
    m = 4 + len(plants)
    fs.cell(row=m, column=1, value="Total").font = Font(bold=True)
    c = fs.cell(row=m, column=3, value=F["ask_b_total"])
    c.font = Font(bold=True)
    _head(fs, m + 2, ["Plant", "Composite ended", "Below (-) or above the median (log10)"], 30)
    k = m + 3
    for p in plants:
        for day, dev in p["flagged"]:
            fs.cell(row=k, column=1, value=p["code"])
            fs.cell(row=k, column=2, value=day)
            c = fs.cell(row=k, column=3, value=dev)
            c.number_format = "+0.00;-0.00"
            k += 1
    _note(fs, k + 1, "PMMoV load per person on the historian's inflow over the composite's 24 hours against the median "
                     "of the plant's composites of the 365 days before it (methods note 6.3); the contract laboratory's "
                     "copies per microlitre of extract times %s for copies a litre." % "{:,}".format(
                         int(F["contract_factor"])), 3, 42)
    _widths(fs, [8, 22, 20])
    fs.freeze_panes = "A4"

    nt = wb.create_sheet("Notes")
    nt["A1"] = "Notes"
    nt["A1"].font = Font(bold=True, size=11)
    notes = [
        ("Prepared", "13 April 2028, for the verification desk's rota of May to September 2028; data as of 12 April "
                     "2028."),
        ("Planning basis", "Last winter's samples (composites of May to September 2027) read the way the dashboard "
                           "reads them from 1 May 2028 (methods note v4, section 6)."),
        ("Published", "The dashboard's levels on each sample's N1 load per person; the counts agree with the "
                      "dashboard's export of 19 October 2027."),
        ("Shed load", "Each sample's load with the RNA lost in the sewer restored: exp(k(T) x travel time), k(T) = k20 "
                      "x theta^(T - 20) fitted to the paired study, T the composite's mean inlet temperature from the "
                      "historian, the travel time from the network memo (November 2025)."),
        ("Results", "The laboratory's accepted run of each sample and target; the composites of 23 December 2027 to 6 "
                    "January 2028 from the contract laboratory's report WL-2801, its copies per microlitre of extract "
                    "converted through its method's volumes."),
        ("Rounding", "Counts and hours are whole numbers; the decay rate to three decimals and theta to two; losses "
                     "and deviations to two."),
    ]
    for i, (a, b) in enumerate(notes, 3):
        nt.cell(row=i, column=1, value=a).font = Font(bold=True, size=10)
        c = nt.cell(row=i, column=2, value=b)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        nt.row_dimensions[i].height = 44 if len(b) > 110 else 30
    _widths(nt, [16, 90])
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    wb.save(path)
    postprocess.scrub(path, AUTHOR, "2028-04-13")
    postprocess.finish_office(path, "2028-04-12T21:10:00Z", "2028-04-13T03:25:00Z", ORG)


def chart_png(G):
    F = G["fig"]
    rows = sorted(F["plants"], key=lambda p: p["tt"])
    fig, ax = plt.subplots(figsize=(6.8, 6.2), dpi=200)
    y = list(range(len(rows)))[::-1]
    h = 0.38
    for yy, p in zip(y, rows):
        ax.barh(yy + h / 2, p["published"], height=h, color="#B9C3D0", zorder=2)
        col = "#A03A22" if p["change"] else "#3E5A78"
        ax.barh(yy - h / 2, p["shed"], height=h, color=col, zorder=2)
        if p["change"]:
            ax.text(max(p["published"], p["shed"]) + 0.08, yy, "%+d" % p["change"], va="center", fontsize=7.2,
                    color=col, fontweight="bold")
    ax.set_yticks(y)
    ax.set_yticklabels(["%s %s (%d min)" % (p["code"], p["name"], p["tt"]) for p in rows], fontsize=6.8)
    ax.set_xticks([0, 1, 2, 3])
    ax.set_xlim(0, 3)
    ax.set_xlabel("Alert episodes opening in winter 2027 (grey: published; blue: shed load; red: count moves)",
                  fontsize=7.2)
    ax.set_title(G["chart_title"], fontsize=9.3, loc="left", fontweight="bold")
    ax.grid(axis="x", color="#DDDDDD", lw=0.6, zorder=0)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.tick_params(axis="x", labelsize=7.5)
    fig.text(0.01, 0.01, "Plants in order of population-weighted travel time. Source: SPPHS wastewater dashboard; paired "
                         "upstream study 2026.", fontsize=5.8, color="#555555", va="bottom")
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def _docx(path, G):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    F = G["fig"]
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    hdr = sec.header.paragraphs[0]
    hdr.text = "%s, surveillance team" % ORG
    hdr.runs[0].font.size = Pt(8)
    ftr = sec.footer.paragraphs[0]
    ftr.text = "Verification desk rota, winter 2028    Page "
    ftr.runs[0].font.size = Pt(8)
    run = ftr.add_run()
    run.font.size = Pt(8)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)

    def para(text, bold=False, size=None, after=6, italic=False, color=None):
        p = d.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        p.paragraph_format.space_after = Pt(after)
        return p

    para("PLANNING NOTE: VERIFICATION DESK, MAY TO SEPTEMBER 2028", bold=True, size=9, after=2, color="555555")
    para("To: the surveillance manager    From: %s, wastewater surveillance coordinator    Date: 13 April 2028"
         % P.COORDINATOR[0], size=9, after=10, color="555555")
    para(G["lead"], bold=True, size=13, after=8)
    para(G["paras"][0], after=8)
    d.add_picture(io.BytesIO(chart_png(G)), width=Inches(6.2))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    moved = [p for p in F["plants"] if p["change"]]
    tb = d.add_table(rows=1, cols=5)
    tb.style = "Light Grid Accent 1"
    for j, h in enumerate(["Plant", "Travel time (min)", "Published", "Shed load", "Change"]):
        tb.rows[0].cells[j].text = h
    for p in moved:
        cells = tb.add_row().cells
        for j, v in enumerate(["%s %s" % (p["code"], p["name"]), str(p["tt"]), str(p["published"]), str(p["shed"]),
                               "%+d" % p["change"]]):
            cells[j].text = v
    cells = tb.add_row().cells
    for j, v in enumerate(["All 24 plants", "", str(F["totals"]["published"]), str(F["totals"]["shed"]),
                           "%+d" % (F["totals"]["shed"] - F["totals"]["published"])]):
        cells[j].text = v
    for row in tb.rows:
        for j, cell in enumerate(row.cells):
            for p_ in cell.paragraphs:
                p_.paragraph_format.space_after = Pt(0)
                if j >= 1:
                    p_.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for r_ in p_.runs:
                    r_.font.size = Pt(8.5)
    adw = [p for p in F["plants"] if p["code"] == "ADW"][0]
    fmt = lambda iso: "{d.day} {d:%B}".format(d=__import__("datetime").date.fromisoformat(iso))
    para("The plants whose count moves; the other %d keep theirs. Ardenmouth West's one published episode (opened the "
         "week of %s, closed the week of %s 2027) splits in two on the shed load. Source: the dashboard's export "
         "(published); the workbook, sheet Episodes (shed load)." % (
             len(F["plants"]) - len(moved), fmt(adw["episodes_published"][0][0]), fmt(adw["episodes_published"][0][1])),
         italic=True, size=8, after=10, color="555555")
    for t in G["paras"][1:]:
        para(t, after=8)
    para(G["footnote"], size=8.5, after=6, color="444444")
    para(G["sources"], italic=True, size=8.5, after=0)
    d.core_properties.author = AUTHOR
    d.core_properties.title = G["title"]
    d.core_properties.comments = ""
    d.save(path)
    postprocess.scrub(path, AUTHOR, "2028-04-13")
    postprocess.finish_office(path, "2028-04-12T22:40:00Z", "2028-04-13T04:05:00Z", ORG)


def write(folder, G):
    os.makedirs(folder, exist_ok=True)
    _workbook(os.path.join(folder, XLSX), G)
    _docx(os.path.join(folder, DOCX), G)
