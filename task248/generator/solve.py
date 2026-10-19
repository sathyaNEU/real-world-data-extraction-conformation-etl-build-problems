"""task248 golden, computed from the shipped pack only.

The call: each plant's winter 2027 alert episodes on the published load per person and on the shed load, the shed load
restoring each composite's RNA lost in the sewer at the decay rate the paired upstream study gives for the water's
temperature, first order over the plant's population-weighted travel time, with the composite's own inlet temperature
from the historian. The asks: each plant's hours from 1 April 2027 to 31 March 2028 with no valid inlet flow reading
in the historian, and each plant's composites from 4 October 2027 to 30 March 2028 that the faecal-strength check flags.

Every route the checks and the write-up compare (the ladder's rungs, the variants that must agree, the asks' wrong
paths) is a switch on the same reader, so the golden and its rivals run through one pipeline."""
import datetime as dt
import os
import re
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from pypdf import PdfReader

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import clock
import estim as E

LOG10E = np.log10(np.e)
WINTER = (dt.date(2027, 5, 3), dt.date(2027, 9, 27))
REF = (dt.date(2026, 4, 6), dt.date(2027, 3, 29))
ASK_A = (dt.date(2027, 4, 1), dt.date(2028, 3, 31))
ASK_B = (dt.date(2027, 10, 4), dt.date(2028, 3, 30))


class Pack:
    def __init__(self, target, F):
        self.t, self.F = target, F
        p = lambda k: os.path.join(target, F[k])
        self.reg = pd.read_csv(p("register"))
        self.codes = list(self.reg.plant_code)
        self.pop = dict(zip(self.reg.plant_code, self.reg.population_served))
        self.names = {n.upper(): c for c, n in zip(self.reg.plant_code, self.reg.plant_name)}
        self.tt = self._travel(p("travel"))
        self.lab = pd.read_csv(p("lab"), dtype={"result_gc_per_l": float})
        self.contract = self._contract(p("contract"))
        self.study = pd.read_csv(p("study"))
        self._historian(p("scada"))

    # ------------------------------------------------------------------ readers
    def _travel(self, path):
        text = " ".join(pg.extract_text() for pg in PdfReader(path).pages)
        out = {}
        for code in self.codes:
            m = re.search(r"\b%s\s+(\d+)\s+(\d+)\s+(\d+)\s+(\d+)" % code, text)
            out[code] = int(m.group(2))
        return out

    def _contract(self, path):
        ws = load_workbook(path, read_only=True).active
        rows = list(ws.iter_rows(values_only=True))
        meth = " ".join(str(r[0]) for r in rows[:6] if r and r[0])
        ml = float(re.search(r"(\d+) mL of each composite", meth).group(1))
        rs = float(re.search(r"resuspended in (\d+) mL", meth).group(1))
        al = float(re.search(r"of which (\d+) uL was extracted", meth).group(1))
        ul = float(re.search(r"eluted in (\d+) uL", meth).group(1))
        head = [i for i, r in enumerate(rows) if r and r[0] == "Lab ref"][0]
        recs = []
        for r in rows[head + 1:]:
            if not r or r[0] is None:
                continue
            site = str(r[1]).upper()
            code = self._site(site)
            d = r[2].date() if hasattr(r[2], "date") else r[2]
            recs.append(dict(code=code, d=d, analyte=r[4], per_ul=float(r[5]), flow=r[6]))
        df = pd.DataFrame(recs, columns=["code", "d", "analyte", "per_ul", "flow"])
        # copies in the extract are the aliquot's; the aliquot is its share of the resuspended pellet, which holds the
        # copies of the wastewater concentrated
        self.contract_factor = ul * (rs * 1000.0 / al) / ml * 1000.0
        self.contract_naive = ul / ml * 1000.0          # the aliquot step left out
        df["per_l"] = df.per_ul * self.contract_factor
        return df

    def _site(self, site):
        s = site.replace(" WWTP", "").replace(" NTH", " NORTH").replace(" STH", " SOUTH").strip()
        return self.names[s + " WWTP"]

    def _historian(self, path):
        h0 = pd.read_csv(path, dtype={"inlet_flow_m3h": str, "inlet_temperature_c": float})
        self.h_rows = len(h0)
        self.h_dups = int(h0.duplicated().sum())
        self.raw = self._mark(h0)
        self.h = self._mark(h0.drop_duplicates())

    def _mark(self, h):
        h = h.copy()
        h["ts"] = pd.to_datetime(h.timestamp_local, format="%Y-%m-%d %H:%M")
        h = h.sort_values(["plant_code", "ts"], kind="mergesort").reset_index(drop=True)
        f = pd.to_numeric(h.inlet_flow_m3h, errors="coerce")
        # a held reading: every reading after the first of a run of four or more identical consecutive flow readings
        key = h.inlet_flow_m3h.fillna("")
        new = (key != key.shift()) | (h.plant_code != h.plant_code.shift())
        run = new.cumsum()
        size = run.map(run.value_counts())
        first = new
        held = (size >= 4) & ~first & (key != "")
        h["flow"] = f
        h["held"] = held
        h["valid"] = f.notna() & (f > 0) & ~held
        # each row's composite: the 24 hours to 08:00 local
        h["comp"] = (h.ts - pd.Timedelta(hours=8)).dt.date + dt.timedelta(days=1)
        return h

    # ------------------------------------------------------------------ the samples
    def composites(self, runs="accepted", contract=True, contract_factor=None):
        """one row per plant and composite: N1 and PMMoV concentrations (gene copies a litre) and the form's flow"""
        lab = self.lab.copy()
        lab["d"] = pd.to_datetime(lab.composite_end).dt.date
        if runs == "accepted":
            lab = lab[lab.status == "accepted"]
        elif runs == "first":
            lab = lab.sort_values("analysed", kind="mergesort").groupby(["sample_id", "target"]).head(1)
        elif runs == "latest":
            lab = lab.sort_values("analysed", kind="mergesort").groupby(["sample_id", "target"]).tail(1)
        elif runs == "mean":
            lab = lab.groupby(["sample_id", "plant_code", "d", "target", "composite_flow_m3"], as_index=False) \
                .result_gc_per_l.mean()
        piv = lab.pivot_table(index=["plant_code", "d"], columns="target", values="result_gc_per_l", aggfunc="first")
        flow = lab.groupby(["plant_code", "d"]).composite_flow_m3.first()
        df = piv.join(flow).reset_index().rename(columns={"composite_flow_m3": "form"})
        if contract and len(self.contract):
            c = self.contract.copy()
            fac = self.contract_factor if contract_factor is None else contract_factor
            c["v"] = c.per_ul * fac
            cp = c.pivot_table(index=["code", "d"], columns="analyte", values="v", aggfunc="first").reset_index()
            cf = c.groupby(["code", "d"]).flow.first().reset_index()
            cp = cp.merge(cf, on=["code", "d"])
            cp = cp.rename(columns={"code": "plant_code", "SARS-CoV-2 N1": "N1", "flow": "form"})
            df = pd.concat([df, cp[["plant_code", "d", "N1", "PMMoV", "form"]]], ignore_index=True)
        return df.sort_values(["plant_code", "d"]).reset_index(drop=True)

    def window(self, fill="mean"):
        """each plant and composite day: the historian's inflow over the composite's 24 hours (cubic metres) and the
        mean inlet temperature; fill: 'mean' (valid hours' mean times 24), 'valid' (sum of valid hours), 'raw' (every
        row's flow as given, a blank as nothing), 'interp' (gaps interpolated in time)"""
        h = self.h
        if fill == "interp":
            g = []
            for code, x in h.groupby("plant_code", sort=False):
                v = x.flow.where(x.valid)
                v.index = x.ts
                v = v.interpolate(method="time", limit_direction="both")
                g.append(pd.DataFrame({"plant_code": code, "comp": x.comp.values, "q": v.values}))
            g = pd.concat(g)
            agg = g.groupby(["plant_code", "comp"]).q.agg(["sum", "count"])
            vol = agg["sum"] * 24.0 / agg["count"]
        elif fill == "raw":
            vol = h.groupby(["plant_code", "comp"]).flow.sum()
        else:
            v = h[h.valid]
            agg = v.groupby(["plant_code", "comp"]).flow.agg(["sum", "count"])
            vol = agg["sum"] * 24.0 / agg["count"] if fill == "mean" else agg["sum"]
        temp = h.groupby(["plant_code", "comp"]).inlet_temperature_c.mean()
        return vol, temp

    # ------------------------------------------------------------------ the decay fits on the paired study
    def pairs(self):
        st = self.study.copy()
        g = st.groupby(["plant_code", "visit_date", "point"]).agg(n1=("n1_copies_per_ml", "mean"),
                                                                    pm=("pmmov_copies_per_ml", "mean"),
                                                                    T=("water_temperature_c", "first")).reset_index()
        up = g[g.point == "upstream"].set_index(["plant_code", "visit_date"])
        pl = g[g.point == "plant"].set_index(["plant_code", "visit_date"])
        j = up.join(pl, lsuffix="_up", rsuffix="_pl").reset_index()
        j["tt"] = j.plant_code.map(self.tt) / 60.0
        j["y"] = np.log(j.n1_pl / j.n1_up)
        j["ypm"] = np.log((j.n1_pl / j.pm_pl) / (j.n1_up / j.pm_up))
        j["T"] = (j.T_up + j.T_pl) / 2.0
        j["month"] = pd.to_datetime(j.visit_date).dt.month
        return j

    def fit(self, form="arrhenius", temp="T", target="y"):
        """k(T) per hour from ln(plant / upstream) = -k(T) tt"""
        j = self.pairs()
        t, T, y = j.tt.values, j[temp].values, j[target].values
        if form == "arrhenius":
            best = None
            for th in np.linspace(1.0, 1.2, 2001):
                g = t * th ** (T - 20)
                k20 = -np.sum(g * y) / np.sum(g * g)
                r = np.sum((y + k20 * g) ** 2)
                if best is None or r < best[0]:
                    best = (r, k20, th)
            _, k20, th = best
            return (lambda TT: k20 * th ** (np.asarray(TT, dtype=float) - 20)), dict(k20=k20, theta=th)
        if form == "linear":
            a, b = np.linalg.lstsq(np.c_[t, t * T], -y, rcond=None)[0]
            return (lambda TT: a + b * np.asarray(TT, dtype=float)), dict(a=a, b=b)
        if form == "per_plant":
            ab = []
            for code, x in j.groupby("plant_code"):
                A = np.c_[np.ones(len(x)), x[temp].values]
                ab.append(np.linalg.lstsq(A, -x[target].values / x.tt.values, rcond=None)[0])
            a, b = np.mean(ab, axis=0)
            return (lambda TT: a + b * np.asarray(TT, dtype=float)), dict(a=a, b=b)
        if form == "constant":
            k = -np.sum(t * y) / np.sum(t * t)
            return (lambda TT: k + 0 * np.asarray(TT, dtype=float)), dict(k=k)
        raise ValueError(form)

    def additive(self):
        """ln(plant / upstream) = c0 + c1 tt + c2 T: a temperature term beside the travel time"""
        j = self.pairs()
        c = np.linalg.lstsq(np.c_[np.ones(len(j)), j.tt.values, j["T"].values], j.y.values, rcond=None)[0]
        return c

    # ------------------------------------------------------------------ the call
    def loads(self, flow="form", runs="accepted"):
        """each plant's composites with log10 N1 and PMMoV loads per person per day, the window's temperature"""
        c = self.composites(runs=runs)
        vol, temp = self.window()
        key = list(zip(c.plant_code, c.d))
        c["hflow"] = [vol.get(k, np.nan) for k in key]
        c["T"] = [temp.get(k, np.nan) for k in key]
        q = c.form if flow == "form" else c.hflow
        c["pop"] = c.plant_code.map(self.pop)
        c["L"] = np.log10(c.N1 * q * 1000.0 / c["pop"])
        c["LP"] = np.log10(c.PMMoV * q * 1000.0 / c["pop"])
        c["tt"] = c.plant_code.map(self.tt) / 60.0
        return c

    def counts(self, corr=None, flow="form", runs="accepted", value="L", restrict=None):
        """each plant's winter 2027 episodes on log10 loads plus corr(row) (log10), and its levels"""
        c = self.loads(flow=flow, runs=runs)
        out, levs = {}, {}
        for code, x in c.groupby("plant_code", sort=False):
            v = x[value].values.copy()
            if corr is not None and (restrict is None or code in restrict):
                v = v + np.array([corr(r) for r in x.itertuples()])
            lev, base = E.levels(list(x.d), list(v), REF)
            out[code] = E.episodes(lev, WINTER)[0]
            levs[code] = (lev, base)
        return out, levs

    def shed_corr(self, form="arrhenius", temp="T", target="y"):
        k, par = self.fit(form, temp, target)
        return (lambda r: LOG10E * float(k(r.T)) * r.tt), par

    def calendar_corr(self, by="month"):
        """the fitted rate applied with each month's (or season's) mean inlet temperature at the plant"""
        k, _ = self.fit()
        c = self.loads()
        key = c.d.map(lambda d: d.month if by == "month" else {12: 0, 1: 0, 2: 0, 3: 1, 4: 1, 5: 1, 6: 2, 7: 2, 8: 2,
                                                                  9: 3, 10: 3, 11: 3}[d.month])
        c["key"] = key.values
        m = c.groupby(["plant_code", "key"]).T.mean()
        return lambda r: LOG10E * float(k(m[(r.plant_code, r.d.month if by == "month" else
                                             {12: 0, 1: 0, 2: 0, 3: 1, 4: 1, 5: 1, 6: 2, 7: 2, 8: 2, 9: 3, 10: 3,
                                              11: 3}[r.d.month])])) * r.tt

    def rungs(self):
        """the ladder: every route's per-plant counts"""
        R = {}
        R["published"] = self.counts()[0]
        corr, par = self.shed_corr()
        R["shed"] = self.counts(corr)[0]
        kc, _ = self.fit("constant")
        R["constant"] = self.counts(lambda r: LOG10E * float(kc(r.T)) * r.tt)[0]
        R["pmmov"] = self.counts(lambda r: -r.LP, value="L")[0]
        c = self.additive()
        R["additive"] = self.counts(lambda r: -LOG10E * (c[1] * r.tt + c[2] * r.T))[0]
        R["study plants only"] = self.counts(corr, restrict={"DNK", "NTL", "ADS", "CLF"})[0]
        return R, par

    def variants(self):
        """the routes that must all give the golden: every temperature-dependent fit, both temperatures, both flows,
        the calendar temperatures, the marker-normalized ratio"""
        V = {}
        for form in ("arrhenius", "linear", "per_plant"):
            for temp in ("T", "T_up", "T_pl"):
                V["%s/%s" % (form, temp)] = self.counts(self.shed_corr(form, temp)[0])[0]
        V["arrhenius/historian flow"] = self.counts(self.shed_corr()[0], flow="hist")[0]
        V["arrhenius/N1 over PMMoV in the study"] = self.counts(self.shed_corr(target="ypm")[0])[0]
        V["arrhenius/monthly temperature"] = self.counts(self.calendar_corr("month"))[0]
        V["arrhenius/seasonal temperature"] = self.counts(self.calendar_corr("season"))[0]
        V["arrhenius/temperature over the travel time before the composite"] = self.counts(self.lagged_corr())[0]
        return V

    def lagged_corr(self):
        """the fitted rate at the inlet temperature of the 24 hours ending one travel time before the composite ended,
        when the composite's sewage was on its way"""
        k, _ = self.fit()
        by = {c: g.set_index("ts").inlet_temperature_c.sort_index() for c, g in self.h.groupby("plant_code")}

        def corr(r):
            end = pd.Timestamp(dt.datetime.combine(r.d, dt.time(8))) - pd.Timedelta(hours=r.tt)
            s = by[r.plant_code]
            T = float(s[(s.index >= end - pd.Timedelta(hours=24)) & (s.index < end)].mean())
            return LOG10E * float(k(T)) * r.tt
        return corr

    # ------------------------------------------------------------------ Ask A
    def hours_window(self):
        """the local hour labels of the window, from the start of 1 April 2027 to the end of 31 March 2028, one per real
        hour (the hour the clock skips in September has none, the hour it repeats in April has two)"""
        u0 = clock.to_utc(dt.datetime.combine(ASK_A[0], dt.time(0)))
        u1 = clock.to_utc(dt.datetime.combine(ASK_A[1] + dt.timedelta(days=1), dt.time(0)))
        n = int((u1 - u0).total_seconds() // 3600)
        return [clock.to_local(u0 + dt.timedelta(hours=k)) for k in range(n)]

    def ask_a(self, grid="local", holds="after_first", absent=True, dedup="exact"):
        """grid 'local' (the clock's real hours) or 'naive' (every wall-clock hour of every day); holds 'after_first'
        or 'all' (every reading of the run); absent False counts only the rows the export has; dedup 'exact' (rows equal
        in every field), 'key' (one row per plant and timestamp) or 'none'"""
        import collections
        h = self.raw if dedup != "exact" else self.h
        if dedup == "key":
            h = h.drop_duplicates(["plant_code", "timestamp_local"])
        lo = pd.Timestamp(dt.datetime.combine(ASK_A[0], dt.time(0)))
        hi = pd.Timestamp(dt.datetime.combine(ASK_A[1], dt.time(23)))
        w = h[(h.ts >= lo) & (h.ts <= hi)]
        valid = w.valid.copy()
        if holds == "all":
            key = w.inlet_flow_m3h.fillna("")
            new = (key != key.shift()) | (w.plant_code != w.plant_code.shift())
            run = new.cumsum()
            size = run.map(run.value_counts())
            valid = valid & ~((size >= 4) & (key != ""))
        if grid == "local":
            need = collections.Counter(self.hours_window())
        else:
            need = collections.Counter(pd.date_range(lo, hi, freq="h").to_pydatetime())
        out = {}
        for code in self.codes:
            x = w[w.plant_code == code]
            if not absent:
                out[code] = int((~valid[x.index]).sum())
                continue
            have = collections.Counter(x.ts[valid[x.index]].dt.to_pydatetime())
            out[code] = sum(max(0, n - have.get(t, 0)) for t, n in need.items())
        return out

    # ------------------------------------------------------------------ Ask B
    def ask_b(self, flow="historian", fill="mean", runs="accepted", contract=True, contract_factor=None,
              sides=2, threshold=0.5):
        c = self.composites(runs=runs, contract=contract, contract_factor=contract_factor)
        if flow == "historian":
            vol, _ = self.window(fill)
            q = np.array([vol.get(k, np.nan) for k in zip(c.plant_code, c.d)])
        else:
            q = c.form.values
        c["LP"] = np.log10(c.PMMoV.values * q * 1000.0 / c.plant_code.map(self.pop).values)
        out, detail = {}, {}
        for code, x in c.groupby("plant_code", sort=False):
            d = list(x.d)
            v = list(x.LP)
            n, flagged = 0, []
            for i, di in enumerate(d):
                if not (ASK_B[0] <= di <= ASK_B[1]):
                    continue
                prev = [v[j] for j in range(len(d)) if di - dt.timedelta(days=365) <= d[j] < di and np.isfinite(v[j])]
                med = float(np.median(prev))
                dev = v[i] - med
                hit = abs(dev) > threshold if sides == 2 else dev < -threshold
                if hit:
                    n += 1
                    flagged.append((di, round(dev, 3)))
            out[code] = n
            detail[code] = flagged
        return out, detail


def solve(target, F):
    pk = Pack(target, F)
    pub, plev = pk.counts()
    corr, par = pk.shed_corr()
    shed, slev = pk.counts(corr)
    rows = []
    for code in pk.codes:
        rows.append(dict(plant=code, name=pk.reg.set_index("plant_code").plant_name[code], tt=pk.tt[code],
                         published=pub[code], shed=shed[code], change=shed[code] - pub[code]))
    return pk, par, rows
