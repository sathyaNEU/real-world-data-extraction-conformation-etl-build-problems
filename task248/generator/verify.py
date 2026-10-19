"""task248 independent verifier: recomputes every graded figure from the shipped bundle alone (no generator module is
imported, the clocks come from the IANA zone, the decay fit from a non-linear least squares) and checks figures.json,
submission.md and the two goldens against it.

    python3 task248/generator/verify.py [target golden figures.json submission.md]
"""
import datetime as dt
import json
import os
import re
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from pypdf import PdfReader
from scipy.optimize import curve_fit

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
F = {"register": "wastewater_sites_register.csv", "lab": "phl_wastewater_results_2026-01_2028-04.csv",
     "contract": "rangeview_report_WL-2801.xlsx", "scada": "spw_inlet_historian_hourly_2026-01-05_2028-04-09.csv",
     "travel": "spw_travel_times_memo_2025-11.pdf", "study": "paired_upstream_study_2026_results.csv",
     "dashboard": "dashboard_export_winter_2027.xlsx", "methods": "wastewater_dashboard_methods_v4.pdf",
     "scadanote": "spw_historian_export_note.pdf", "labnote": "phl_results_export_notes.pdf"}
TZ = "Pacific/Auckland"
REF = (pd.Timestamp("2026-04-06"), pd.Timestamp("2027-03-29"))
WIN = (pd.Timestamp("2027-05-03"), pd.Timestamp("2027-09-27"))


def fail(msg):
    print("VERIFY FAIL:", msg)
    sys.exit(1)


def check(cond, msg):
    if not cond:
        fail(msg)


def pdf_text(path):
    return " ".join(" ".join(p.extract_text() for p in PdfReader(path).pages).split())


def episodes(levels):
    """levels: Series indexed by Monday; opens at 1.0 or more, closes below 0.5; counted when it opens in the winter"""
    n, open_ = 0, False
    for w, x in levels.sort_index().items():
        if not open_ and x >= 1.0:
            open_ = True
            n += WIN[0] <= w <= WIN[1]
        elif open_ and x < 0.5:
            open_ = False
    return n


def plant_levels(df, col):
    """df: plant, end (Timestamp), col (log10 load); weekly by the Monday of the composite's end day"""
    df = df.assign(week=df.end.dt.to_period("W-SUN").dt.start_time)
    out = {}
    for code, x in df.groupby("plant"):
        wk = x.groupby("week")[col].mean()
        base = np.percentile(wk[(wk.index >= REF[0]) & (wk.index <= REF[1])].values, 10)
        out[code] = wk - base
    return out


def load(target):
    reg = pd.read_csv(os.path.join(target, F["register"]))
    pop = dict(zip(reg.plant_code, reg.population_served))
    txt = pdf_text(os.path.join(target, F["travel"]))
    tt = {c: int(re.search(r"%s (\d+) (\d+) (\d+) (\d+)" % c, txt).group(2)) for c in reg.plant_code}
    lab = pd.read_csv(os.path.join(target, F["lab"]))
    lab = lab[lab.status == "accepted"]
    lab["end"] = pd.to_datetime(lab.composite_end).dt.normalize()
    piv = lab.pivot_table(index=["plant_code", "end"], columns="target", values="result_gc_per_l", aggfunc="first")
    piv["form"] = lab.groupby(["plant_code", "end"]).composite_flow_m3.first()
    S = piv.reset_index().rename(columns={"plant_code": "plant"})
    # the contract laboratory's report: copies per microlitre of extract to copies a litre through its method's volumes
    ws = load_workbook(os.path.join(target, F["contract"]), read_only=True).active
    allrows = list(ws.iter_rows(values_only=True))
    meth = " ".join(str(r[0]) for r in allrows[:7] if r and r[0])
    v = [float(x) for x in re.findall(r"(\d+) (?:mL|uL)", meth)]
    sample_ml, resusp_ml, aliquot_ul, elute_ul = v[:4]
    factor = elute_ul / aliquot_ul * (resusp_ml * 1000.0) / sample_ml * 1000.0
    names = {n.upper().replace(" WWTP", ""): c for c, n in zip(reg.plant_code, reg.plant_name)}
    alias = {"ARDENMOUTH NTH": "ARDENMOUTH NORTH", "ARDENMOUTH STH": "ARDENMOUTH SOUTH"}
    rec = []
    start = [i for i, r in enumerate(allrows) if r and r[0] == "Lab ref"][0] + 1
    for r in allrows[start:]:
        if not r or not r[0]:
            continue
        site = str(r[1]).upper().replace(" WWTP", "")
        rec.append(dict(plant=names[alias.get(site, site)], end=pd.Timestamp(r[2]).normalize(),
                        target="N1" if "N1" in r[4] else "PMMoV", v=float(r[5]) * factor, form=r[6]))
    c = pd.DataFrame(rec).pivot_table(index=["plant", "end"], columns="target", values="v", aggfunc="first")
    c["form"] = pd.DataFrame(rec).groupby(["plant", "end"]).form.first()
    S = pd.concat([S, c.reset_index()], ignore_index=True).sort_values(["plant", "end"]).reset_index(drop=True)
    return reg, pop, tt, S, factor


def historian(target):
    h = pd.read_csv(os.path.join(target, F["scada"]), dtype={"inlet_flow_m3h": str}).drop_duplicates()
    parts = []
    for code, x in h.groupby("plant_code", sort=False):
        x = x.copy()
        loc = pd.to_datetime(x.timestamp_local).dt.tz_localize(TZ, ambiguous="infer")
        x["utc"] = loc.dt.tz_convert("UTC")
        f = pd.to_numeric(x.inlet_flow_m3h, errors="coerce")
        s = x.inlet_flow_m3h.fillna("#")
        grp = (s != s.shift()).cumsum()
        n = grp.map(grp.value_counts())
        pos = x.groupby(grp).cumcount()
        held = (n >= 4) & (pos > 0) & (s != "#")
        x["ok"] = f.notna() & (f > 0) & ~held
        x["f"] = f
        parts.append(x)
    return pd.concat(parts, ignore_index=True)


def windows(H, S):
    """each composite: the mean inlet temperature and the inflow (valid hours' mean times 24) over the 24 hours to
    08:00 local on its day"""
    end = S.end.dt.tz_localize(None) + pd.Timedelta(hours=8)
    end_utc = end.dt.tz_localize(TZ).dt.tz_convert("UTC")
    S = S.assign(e=end_utc, s=end_utc - pd.Timedelta(hours=24))
    T, Q = [], []
    by = {k: g.set_index("utc").sort_index() for k, g in H.groupby("plant_code")}
    for r in S.itertuples():
        g = by[r.plant]
        w = g[(g.index >= r.s) & (g.index < r.e)]
        T.append(w.inlet_temperature_c.mean())
        Q.append(w.f[w.ok].mean() * 24.0)
    return S.assign(T=T, Q=Q)


def decay(target, tt):
    st = pd.read_csv(os.path.join(target, F["study"]))
    m = st.groupby(["plant_code", "visit_date", "point"]).agg(n1=("n1_copies_per_ml", "mean"),
                                                                temp=("water_temperature_c", "mean")).unstack("point")
    y = np.log(m[("n1", "plant")] / m[("n1", "upstream")]).values
    T = ((m[("temp", "plant")] + m[("temp", "upstream")]) / 2).values
    t = np.array([tt[c] / 60.0 for c in m.index.get_level_values(0)])
    (k20, th), _ = curve_fit(lambda X, a, b: -a * b ** (X[1] - 20) * X[0], (t, T), y, p0=(0.04, 1.05))
    return k20, th, len(y), m, t


def main(target, golden, figpath, subpath):
    fig = json.load(open(figpath))
    sub = open(subpath, encoding="utf-8").read()
    reg, pop, tt, S, factor = load(target)
    H = historian(target)
    W = windows(H, S)
    W["pop"] = W.plant.map(pop)
    W["L"] = np.log10(W.N1 * W.form * 1000.0 / W["pop"])
    k20, th, nvis, m, t = decay(target, tt)
    W["Ls"] = W.L + np.log10(np.e) * k20 * th ** (W["T"] - 20) * W.plant.map(tt) / 60.0
    W["Lh"] = np.log10(W.N1 * W.Q * 1000.0 / W["pop"]) + np.log10(np.e) * k20 * th ** (W["T"] - 20) * \
        W.plant.map(tt) / 60.0
    main_rows = W[W.end <= pd.Timestamp("2027-10-03")]
    pub = {c: episodes(v) for c, v in plant_levels(main_rows, "L").items()}
    shed = {c: episodes(v) for c, v in plant_levels(main_rows, "Ls").items()}
    shed_h = {c: episodes(v) for c, v in plant_levels(main_rows, "Lh").items()}
    n = 0
    check(shed_h == shed, ("the historian's flow", shed_h))
    check(sum(pub.values()) == fig["totals"]["published"] == 30, ("published", pub))
    check(sum(shed.values()) == fig["totals"]["shed"] == fig["call"] == 26, ("shed", shed))
    moves = {c: shed[c] - pub[c] for c in reg.plant_code if shed[c] != pub[c]}
    check(moves == fig["moves"], ("moves", moves))
    P_ = {p["code"]: p for p in fig["plants"]}
    for c in reg.plant_code:
        check(P_[c]["published"] == pub[c] and P_[c]["shed"] == shed[c] and P_[c]["tt"] == tt[c], ("plant", c))
    check("%.3f" % k20 == fig["decay"]["k20"] and "%.2f" % th == fig["decay"]["theta"] and
          int(round((th - 1) * 100)) == fig["decay"]["theta_pct"] and nvis == fig["decay"]["visits"],
          ("the decay fit", k20, th, nvis))
    n += 6 + len(reg)
    # the dashboard's export is the published basis
    ep = pd.read_excel(os.path.join(target, F["dashboard"]), sheet_name="Alert episodes", header=2)
    check(ep.groupby("Plant").size().to_dict() == {c: v for c, v in pub.items() if v}, "the dashboard's episodes")
    # the study's seasonal losses
    loss = -np.log10(m[("n1", "plant")] / m[("n1", "upstream")])
    mon = pd.to_datetime(pd.Series(m.index.get_level_values(1))).dt.month.values
    for code in ("CLF", "ADS"):
        sel = m.index.get_level_values(0) == code
        for name, months in (("December to February", (12, 1, 2)), ("June to August", (6, 7, 8))):
            v = round(float(loss[sel & np.isin(mon, months)].mean()), 2)
            check(v == fig["seasonal"][code][name], ("seasonal loss", code, name, v))
            check("%.2f" % v in sub, ("the submission's seasonal loss", code, name))
            n += 2
    # a constant rate gives back the published counts
    kc = -np.sum(t * np.log(m[("n1", "plant")] / m[("n1", "upstream")]).values) / np.sum(t * t)
    W["Lc"] = W.L + np.log10(np.e) * kc * W.plant.map(tt) / 60.0
    const = {c: episodes(v) for c, v in plant_levels(W[W.end <= pd.Timestamp("2027-10-03")], "Lc").items()}
    check(const == pub and fig["rungs"]["constant"] == 30, "a constant rate")
    n += 2
    # ask A: real hours of the year (UTC) without a valid reading
    u0 = pd.Timestamp("2027-04-01 00:00").tz_localize(TZ).tz_convert("UTC")
    u1 = pd.Timestamp("2028-04-01 00:00").tz_localize(TZ).tz_convert("UTC")
    hours = pd.date_range(u0, u1, freq="h", inclusive="left")
    check(len(hours) == 8784 == fig["conventions"]["real_hours"], "8,784 hours")
    A = {}
    for code, x in H.groupby("plant_code"):
        okh = set(x.utc[x.ok])
        A[code] = int(sum(1 for u in hours if u not in okh))
    for c in reg.plant_code:
        check(A[c] == P_[c]["hours"], ("ask A", c, A[c], P_[c]["hours"]))
        check("   - %s %s: %d" % (c, P_[c]["name"], A[c]) in sub, ("ask A line", c))
        d = P_[c]["hours_detail"]
        check(d["no_row"] + d["blank"] + d["held"] + d["zero_or_below"] == A[c], ("ask A detail", c))
        n += 3
    # ask B: the faecal-strength check
    W["LP"] = np.log10(W.PMMoV * W.Q * 1000.0 / W["pop"])
    B = {}
    for code, x in W.groupby("plant"):
        x = x.set_index("end").sort_index()
        cnt = 0
        for e, r in x.iterrows():
            if not (pd.Timestamp("2027-10-04") <= e <= pd.Timestamp("2028-03-30")):
                continue
            prev = x.LP[(x.index >= e - pd.Timedelta(days=365)) & (x.index < e)]
            cnt += abs(r.LP - prev.median()) > 0.5
        B[code] = int(cnt)
    for c in reg.plant_code:
        check(B[c] == P_[c]["flags"], ("ask B", c, B[c], P_[c]["flags"]))
        check("   - %s %s: %d" % (c, P_[c]["name"], B[c]) in sub, ("ask B line", c))
        n += 2
    check(abs(factor - fig["contract_factor"]) < 1e-9 and "{:,}".format(int(factor)) in sub, "the contract factor")
    # the goldens carry the same figures
    wb = load_workbook(os.path.join(golden, "winter_2027_episodes_shed_load.xlsx"), read_only=True, data_only=True)
    got = {r[0]: r for r in wb["Episodes"].iter_rows(min_row=5, values_only=True) if r[0] in pub}
    tot = [r for r in wb["Episodes"].iter_rows(min_row=5, values_only=True) if r[0] == "Total"][0]
    check(tot[3] == 30 and tot[4] == 26, "the golden's totals")
    for c in reg.plant_code:
        check(got[c][3] == pub[c] and got[c][4] == shed[c], ("golden row", c))
    a = {r[0]: r for r in wb["Historian hours"].iter_rows(min_row=4, values_only=True) if r[0] in pub}
    b = {r[0]: r for r in wb["Faecal-strength check"].iter_rows(min_row=4, max_row=27, values_only=True) if r[0] in pub}
    for c in reg.plant_code:
        check(a[c][6] == A[c] and b[c][2] == B[c], ("golden asks", c))
    n += 3 + 2 * len(reg)
    import docx
    doc = docx.Document(os.path.join(golden, "verification_desk_winter_2028.docx"))
    text = " ".join(p.text for p in doc.paragraphs)
    check("26 alert episodes" in text and len(doc.inline_shapes) == 1, "the note's number and chart")
    for c in moves:
        check(c in " ".join(cell.text for t in doc.tables for row in t.rows for cell in row.cells), ("the note's table", c))
    n += 1 + len(moves)
    # the filed sentences the call and the asks lean on
    docs = {F["methods"]: ["the laboratory's accepted analysis", "of the day its composite ended",
                           "(PERCENTILE.INC)", "the weeks starting 6 April 2026 to 29 March 2027",
                           "first week a plant's level is at High or above", "back at Low",
                           "the weeks starting 3 May to 27 September 2027", "the RNA the sewer destroyed",
                           "exactly as sections 4 and 5 compute them", "more than 0.5 log10 above or below",
                           "in the 365 days before the composite's day", "from the historian"],
            F["scadanote"]: ["four or more identical consecutive flow readings", "a flow of zero or below is a meter fault",
                             "A blank flow", "New Zealand local time, the start of each hour",
                             "23:00 on 9 April 2028, %s hours, with a row" % "{:,}".format(
                                 len(pd.date_range(pd.Timestamp("2026-01-05").tz_localize(TZ),
                                                   pd.Timestamp("2028-04-10").tz_localize(TZ), freq="h",
                                                   inclusive="left")))],
            F["labnote"]: ["The accepted run is the sample's result for the target"],
            "paired_upstream_study_2026_field_report.pdf": ["Twenty-six visits were planned at each plant",
                                                            "Netherleigh on 8 June 2026", "gene copies per millilitre"]}
    for f, phrases in docs.items():
        t_ = pdf_text(os.path.join(target, f))
        for ph in phrases:
            check(ph in t_, ("filed sentence", f, ph))
            n += 1
    eml = " ".join(open(os.path.join(target, "desk_rota_winter_2028_request.eml"), encoding="utf-8").read().split())
    check("1 April 2027 to 31 March 2028" in eml and "4 October 2027 to 30 March 2028" in eml, "the request's windows")
    dct = " ".join(open(os.path.join(target, "DATA_DICTIONARY.md"), encoding="utf-8").read().split())
    check("(gene copies per millilitre)" in dct and "the start of the hour" in dct, "the dictionary")
    cont = open(os.path.join(target, "folder_contents.md"), encoding="utf-8").read()
    listed = [f for f in os.listdir(target) if f != "folder_contents.md"]
    check(all(f in cont for f in listed), ("the folder's contents list every file", [f for f in listed if f not in cont]))
    check(chr(0x2014) not in sub and chr(0x2013) not in sub, "no dashes")
    n += 4
    print("verify: %d checks passed" % n)
    return n


if __name__ == "__main__":
    a = sys.argv[1:]
    if len(a) == 4:
        main(*a)
    else:
        main(os.path.join(TASK, "target"), os.path.join(TASK, "golden"), os.path.join(HERE, "figures.json"),
             os.path.join(TASK, "submission.md"))
