"""The pack's documents and workbooks: the district's inventory method, the 2027 grant round's guidelines, the receiver
network's notes on its export, the district's notes on its FAA register extract, the four count studies, the state's
based aircraft validation and airport directory, and the airports' 5010 records."""
import datetime as dt

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import letter                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import inch                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import pack as P      # noqa: E402
import world as W     # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("LiberationSerif", "Regular"), ("LiberationSerif-Bold", "Bold"), ("LiberationSerif-Italic", "Italic"),
              ("LiberationSans", "Regular"), ("LiberationSans-Bold", "Bold"), ("LiberationSans-Italic", "Italic")):
    try:
        pdfmetrics.getFont(nm)
    except KeyError:
        pdfmetrics.registerFont(TTFont(nm, FONT_DIR + nm.split("-")[0] + "-" + f + ".ttf"))
INK = colors.HexColor("#1c1e21")
GREY = colors.HexColor("#5a5f66")
GREEN = colors.HexColor("#28543c")


def styles(font="LiberationSans", accent=GREEN, size=9.8):
    bold = font + "-Bold"
    return dict(
        BODY=ParagraphStyle("body", fontName=font, fontSize=size, leading=size * 1.37, textColor=INK, spaceAfter=6),
        CELL=ParagraphStyle("cell", fontName=font, fontSize=8.2, leading=10.2, textColor=INK),
        CELLB=ParagraphStyle("cellb", fontName=bold, fontSize=8.2, leading=10.2, textColor=INK),
        H1=ParagraphStyle("h1", fontName=bold, fontSize=14, leading=18, textColor=accent, spaceAfter=2),
        H1S=ParagraphStyle("h1s", fontName=font, fontSize=9, leading=12, textColor=GREY, spaceAfter=10),
        H2=ParagraphStyle("h2", fontName=bold, fontSize=10.6, leading=14, textColor=INK, spaceBefore=8, spaceAfter=3))


S = styles()


def _doc(path, title, footer, author):
    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(GREY)
        canvas.setFont("LiberationSans", 7.5)
        canvas.drawString(0.9 * inch, 0.6 * inch, footer)
        canvas.drawRightString(letter[0] - 0.9 * inch, 0.6 * inch, "Page %d" % doc.page)
        canvas.restoreState()
    d = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.9 * inch, rightMargin=0.9 * inch, topMargin=0.8 * inch,
                          bottomMargin=0.9 * inch, title=title, author=author, subject="", creator="", producer="")
    return d, frame


def _table(rows, widths, st=S, head="#e2ebe4", align_right_from=None):
    data = [[Paragraph(str(v), st["CELLB"] if i == 0 else st["CELL"]) for v in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    style = [("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa0a6")),
             ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(head)),
             ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 4),
             ("RIGHTPADDING", (0, 0), (-1, -1), 4)]
    t.setStyle(TableStyle(style))
    return t


def _story(title, sub, sections, st=S):
    story = [Paragraph(title, st["H1"]), Paragraph(sub, st["H1S"])]
    for head, paras in sections:
        if head:
            story.append(Paragraph(head, st["H2"]))
        for p in paras:
            story.append(p if not isinstance(p, str) else Paragraph(p, st["BODY"]))
    return story


# ---------------------------------------------------------------------------------------------- the method
def lead_table():
    L = W.LEAD
    return _table([["Engines", "Takeoff", "Landing", "Each operation of a touch-and-go"],
                   ["One (single-engine)", "%.2f g" % L[("single", "departure")], "%.2f g" % L[("single", "arrival")],
                    "%.2f g" % L[("single", "touch_and_go")]],
                   ["Two or more (multi-engine)", "%.2f g" % L[("multi", "departure")],
                    "%.2f g" % L[("multi", "arrival")], "%.2f g" % L[("multi", "touch_and_go")]]],
                  [2.0 * inch, 1.1 * inch, 1.1 * inch, 2.4 * inch])


def write_method(path):
    d, frame = _doc(path, "%s Revision 3" % W.METHOD_ID,
                    "%s  |  %s, Revision 3" % (W.DISTRICT, W.METHOD_ID), W.DISTRICT)
    secs = [
        ("1. Purpose and scope", [
            "1.1 This method sets how the District inventories the lead emitted by piston-engine aircraft taking off "
            "and landing at the 30 public-use airfields in the Corrisan Valley, field by field and year by year. The "
            "inventory year runs from 1 January to 31 December in local time (Pacific time).",
            "1.2 Lead comes from the tetraethyl lead in aviation gasoline, which only piston engines burn. Aircraft "
            "with turbine engines, gliders and balloons are outside the inventory, and so is lead emitted in flight away "
            "from the field."]),
        ("2. Operations", [
            "2.1 An operation is a takeoff or a landing. A touch-and-go is two operations, a landing and a takeoff, as "
            "a control tower counts it.",
            "2.2 A field's operations in the inventory year are every operation by a piston-engine aircraft at the "
            "field in that year, based or visiting, whether or not the District's receiver network recorded it."]),
        ("3. Lead per operation", [
            "3.1 An operation the network recorded is assigned the lead in Table 1 for its aircraft's number of engines "
            "and the kind of operation.",
            lead_table(), Spacer(1, 6),
            "Table 1. Lead per operation, piston-engine aircraft. Factors from the District's 2023 fuel and "
            "time-in-mode survey of the valley's fleet (Revision 2).",
            "3.2 An operation the network did not record is assigned %.3f g, the mean of a single-engine takeoff and "
            "landing, since the aircraft that flew it is not known." % W.LEAD_UNRECORDED,
            "3.3 A field's lead for the year is the sum over its operations, reported in kilograms."]),
        ("4. Sources", [
            "4.1 The District's ADS-B receiver network, one receiver at each field, records the operations of aircraft "
            "that broadcast ADS-B Out; its export gives each operation's time, field and kind and the aircraft's "
            "24-bit address. The FAA aircraft register gives each aircraft's engine type and number of engines.",
            "4.2 At the six fields with a control tower, the operations the network did not record are the tower's "
            "count of civil operations in the tower's hours of operation less the civil operations the network "
            "recorded at the field in those hours; the inventory takes them as operations by piston-engine "
            "aircraft.",
            "4.3 At the four fields where the District commissioned a count of every operation through 2026 (Fennick "
            "Airpark, Lanthorne Field, Graylow Municipal and Dunmarsh Airport), they are the count's piston-engine "
            "operations less those the network recorded.",
            "4.4 At every other field, the inventory estimates them from the fields with a count."]),
        ("5. Reporting", [
            "5.1 The inventory reports each field's lead in kilograms for the year. The District's grant programs "
            "that qualify fields on aircraft lead use this inventory as their guidelines state."]),
        ("Revision history", [
            _table([["Revision", "Adopted", "Change"],
                    ["1", "14 May 2019", "First issue: operations from the airports' 5010 records and tower counts."],
                    ["2", "21 June 2023", "Lead per operation from the District's fuel and time-in-mode survey; "
                                          "touch-and-go factors added."],
                    ["3", "18 March 2026", "Operations from the receiver network, the tower counts and the District's "
                                           "commissioned counts; Section 4 rewritten."]],
                   [0.8 * inch, 1.1 * inch, 4.7 * inch])]),
    ]
    story = _story("Lead from piston-engine aircraft at the valley's airfields: inventory method",
                   "%s, Planning and Air Toxics. Method %s, Revision 3, adopted 18 March 2026." %
                   (W.DISTRICT, W.METHOD_ID), secs)
    d.build(story, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------------------------------------- the grant round
def write_grants(path):
    d, frame = _doc(path, "Unleaded Avgas Transition Grants, 2027 round",
                    "%s  |  Unleaded Avgas Transition Grants, 2027 round" % W.DISTRICT, W.DISTRICT)
    secs = [
        ("1. Purpose", [
            "1.1 The Unleaded Avgas Transition Grants pay for an unleaded aviation gasoline storage tank and dispenser "
            "at public-use airfields in the District, so that the piston-engine aircraft using them can move off "
            "leaded fuel. The program is funded from the District's air toxics reduction account."]),
        ("2. Who qualifies", [
            "2.1 An eligible field is a public-use airfield in the District without an air traffic control tower.",
            "2.2 An eligible field qualifies for the 2027 round if its 2026 lead from piston-engine aircraft, as the "
            "District's aircraft lead inventory estimates it under method %s, Revision 3, is %.1f kg or more."
            % (W.METHOD_ID, W.THRESHOLD_KG),
            "2.3 The round funds every qualifying field."]),
        ("3. What the grant pays", [
            "3.1 Up to $214,000 a field for a tank of 10,000 to 12,000 gallons, its dispenser, card reader and "
            "containment, and the first load of unleaded fuel. The field's owner provides the site and keeps the "
            "installation for at least ten years.",
            "3.2 Grants are paid against invoices in two installments, at the order of the tank and at "
            "commissioning."]),
        ("4. Timetable", [
            "4.1 Staff publish the qualifying fields with the inventory. The Governing Board adopts the round with the "
            "District's 2027 budget at its meeting of 10 March 2027. Field owners have until 30 June 2027 to accept "
            "the grant and must commission the installation by 30 June 2028."]),
    ]
    story = _story("Unleaded Avgas Transition Grants: 2027 round guidelines",
                   "%s, Grants Program. Issued 4 December 2026 for the Governing Board's adoption of the round." %
                   W.DISTRICT, secs)
    d.build(story, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------------------------------------- the network's notes
def write_network(path):
    d, frame = _doc(path, "ADS-B receiver network: the field events export",
                    "%s  |  Air Monitoring, ADS-B receiver network" % W.DISTRICT, W.PEOPLE["network"][0])
    cols = [["Column", "Contents"],
            ["event_id", "The network database's key for the event."],
            ["site", "The field's location identifier."],
            ["event_time_utc", "When the aircraft met or left the runway, UTC, in ISO 8601 form."],
            ["event", "takeoff, landing or touch_and_go. A touch-and-go is one event."],
            ["icao24", "The aircraft's 24-bit address as broadcast, in lower-case hexadecimal. The FAA register lists "
                       "each registration's address as MODE S CODE HEX."],
            ["flight_id", "The flight identification the aircraft broadcast, as its avionics were set: usually the "
                          "N-number, sometimes blank."],
            ["emitter_cat", "The emitter category broadcast: A1 light (under 15,500 lb), A2 small (15,500 to "
                            "75,000 lb), A7 rotorcraft, B1 glider."],
            ["runway", "The runway end used."]]
    secs = [
        ("The network", [
            "Since spring 2025 the District has run an ADS-B receiver at each of the 30 public-use airfields in the "
            "valley. Each receiver decodes ADS-B Out broadcasts on 1090 MHz (extended squitter) and 978 MHz (UAT) and "
            "turns each track that meets or leaves a runway into an event: a takeoff, a landing or a touch-and-go.",
            "An aircraft that broadcasts ADS-B Out is recorded at every field it uses. An aircraft without ADS-B Out, "
            "or with it switched off, is not recorded. When the receivers were commissioned they were run for four "
            "weeks beside the six control towers' logs, and every takeoff, landing and touch-and-go by an aircraft "
            "broadcasting ADS-B Out was recorded. The export leaves out aircraft broadcasting military addresses."]),
        ("The export", [
            "The export of 8 January 2027 covers 30 December 2025 to 2 January 2027, local time, one row per event.",
            _table(cols, [1.2 * inch, 5.4 * inch]), Spacer(1, 6),
            "All times are UTC. The valley keeps Pacific time: UTC-8, and UTC-7 from 8 March to 1 November 2026 "
            "(02:00 local)."]),
    ]
    story = _story("ADS-B receiver network: notes on the field events export",
                   "%s, Air Monitoring. %s, %s. Revised 8 January 2027." %
                   (W.DISTRICT, W.PEOPLE["network"][0], W.PEOPLE["network"][1]), secs)
    d.build(story, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------------------------------------- the register notes
def write_reg_notes(path):
    d, frame = _doc(path, "FAA register extract", "%s  |  Data Services" % W.DISTRICT, W.DISTRICT)
    secs = [
        ("What the extract holds", [
            "The extract was cut from the FAA's releasable aircraft database as downloaded on 1 February 2027. "
            "MASTER.txt holds the current registration of every aircraft whose 24-bit address the receiver network "
            "recorded from 30 December 2025 to 2 January 2027, of every N-number on the state's 2027 based aircraft "
            "validation for the four counties, and of every other aircraft registered to an address in the four "
            "counties. DEREG.txt holds the cancelled registrations of the same N-numbers and addresses. ACFTREF.txt "
            "and ENGINE.txt hold the reference rows for the aircraft and engine codes the two files use. Owners' names "
            "and addresses are removed under the District's data policy; the files keep the FAA's other columns as "
            "downloaded.",
            "Each file is comma-separated, with every value padded with spaces to its field width and a comma at "
            "the end of each line. N-NUMBER is written without the leading N. Dates are YYYYMMDD. MODE S CODE is the "
            "24-bit address in octal and MODE S CODE HEX the same address in hexadecimal. In DEREG.txt, CANCEL-DATE "
            "is the date the registration was cancelled and EXP-COUNTRY the country an exported aircraft went to."]),
        ("Codes", [
            _table([["Field", "Values"],
                    ["TYPE AIRCRAFT, TYPE-ACFT", "1 glider, 2 balloon, 3 blimp or dirigible, 4 fixed wing single "
                                                 "engine, 5 fixed wing multi engine, 6 rotorcraft, 7 weight-shift "
                                                 "control, 8 powered parachute, 9 gyroplane"],
                    ["TYPE ENGINE, TYPE-ENG, TYPE", "0 none, 1 reciprocating, 2 turbo-prop, 3 turbo-shaft, 4 "
                                                    "turbo-jet, 5 turbo-fan, 6 ramjet, 7 2 cycle, 8 4 cycle, 9 unknown, "
                                                    "10 electric, 11 rotary"],
                    ["TYPE REGISTRANT", "1 individual, 2 partnership, 3 corporation, 4 co-owned, 5 government, 7 "
                                        "LLC, 8 non-citizen corporation, 9 non-citizen co-owned"],
                    ["AC-WEIGHT", "CLASS 1 up to 12,499 lb, CLASS 2 12,500 to 19,999 lb, CLASS 3 20,000 lb and "
                                  "over"],
                    ["BUILD-CERT-IND", "0 type certificated, 1 not type certificated, 2 light sport"],
                    ["NO-ENG", "number of engines"],
                    ["MFR MDL CODE, MFR-MDL-CODE", "the aircraft's code in ACFTREF.txt (CODE)"],
                    ["ENG MFR MDL, ENG-MFR-MDL", "the engine's code in ENGINE.txt (CODE)"]],
                   [1.9 * inch, 4.7 * inch])]),
    ]
    story = _story("Notes on the District's FAA register extract",
                   "%s, Data Services. 3 February 2027." % W.DISTRICT, secs)
    d.build(story, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------------------------------------- count studies
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
          "November", "December"]


def _fmt(n):
    return "{:,}".format(int(n))


def write_f37(path, x):
    st = styles("LiberationSerif", colors.HexColor("#3b3b6d"), 10.2)
    d, frame = _doc(path, "Fennick Airpark (F37) 2026 operations count", "Ardley Counts LLC  |  F37 2026 count",
                    W.STUDY_FIRMS["F37"])
    mon = P.monthly(P.study_counts(x, "F37"))
    rows = [["Month", "Piston-engine airplanes", "Turbine airplanes", "Helicopters", "Gliders", "All operations"]]
    tot = [0, 0, 0, 0, 0]
    for m in range(1, 13):
        v = mon.get(m, {})
        r = [v.get("piston", 0), v.get("turbine", 0), v.get("helicopter", 0), v.get("glider", 0)]
        r.append(sum(r))
        tot = [a + b for a, b in zip(tot, r)]
        rows.append([MONTHS[m - 1]] + [_fmt(z) for z in r])
    rows.append(["2026"] + [_fmt(z) for z in tot])
    secs = [
        ("Summary", [
            "Ardley Counts LLC counted every aircraft operation at Fennick Airpark from 1 January to 31 December 2026 "
            "for the %s. The field recorded %s operations in the year, %s of them by piston-engine airplanes."
            % (W.DISTRICT, _fmt(tot[4]), _fmt(tot[0]))]),
        ("How the count was made", [
            "Two acoustic and video counters, one at each runway end, ran around the clock through the year. An "
            "analyst reviewed every event against its images and classed the aircraft as a "
            "piston-engine airplane, a turbine airplane, a helicopter or a glider. A takeoff or a landing is one "
            "operation and a touch-and-go two. Months are local time."]),
        ("Operations by month", [_table(rows, [1.0 * inch, 1.3 * inch, 1.1 * inch, 0.9 * inch, 0.8 * inch,
                                               1.1 * inch], st, head="#e5e5f0")]),
    ]
    story = _story("Fennick Airpark (F37): 2026 aircraft operations count, final report",
                   "Prepared for the %s by Ardley Counts LLC. 22 January 2027." % W.DISTRICT, secs, st)
    d.build(story, onFirstPage=frame, onLaterPages=frame)


def write_g42(path, x):
    st = styles("LiberationSans", colors.HexColor("#6d3b2a"), 9.6)
    d, frame = _doc(path, "Graylow Municipal Airport 2026 activity count",
                    "Northfork Airport Planning  |  Graylow Municipal Airport (G42)", W.STUDY_FIRMS["G42"])
    mon = P.monthly(P.study_counts(x, "G42"))
    rows = [["Month (2026)", "All aircraft operations", "Piston airplane operations", "Turboprop and jet operations",
             "Rotorcraft operations"]]
    tot = [0, 0, 0, 0]
    for m in range(1, 13):
        v = mon.get(m, {})
        r = [v.get("piston", 0), v.get("turbine", 0), v.get("helicopter", 0)]
        r = [sum(r)] + r
        tot = [a + b for a, b in zip(tot, r)]
        rows.append([MONTHS[m - 1][:3]] + [_fmt(z) for z in r])
    rows.append(["Year"] + [_fmt(z) for z in tot])
    secs = [
        ("1 Scope", [
            "Northfork Airport Planning was engaged by the %s to count aircraft activity at Graylow Municipal Airport "
            "(G42) for calendar 2026, around the clock, by aircraft class." % W.DISTRICT]),
        ("2 Method", [
            "A camera-triggered acoustic counter at midfield logged every runway event; each event was verified from "
            "its images. Operations follow FAA counting practice: a takeoff or a landing is one operation, a "
            "touch-and-go two. Gliders do not use the field."]),
        ("3 Results", [_table(rows, [1.0 * inch, 1.4 * inch, 1.5 * inch, 1.5 * inch, 1.2 * inch], st,
                              head="#f0e6e1"),
                       Spacer(1, 6),
                       "Annual operations: %s, of which piston airplanes %s." % (_fmt(tot[0]), _fmt(tot[1]))]),
    ]
    story = _story("Graylow Municipal Airport (G42) 2026 Activity Count",
                   "Northfork Airport Planning for the %s. Final, 29 January 2027." % W.DISTRICT, secs, st)
    d.build(story, onFirstPage=frame, onLaterPages=frame)


def _head(ws, row, fill="DCE6F1"):
    from openpyxl.styles import Font, PatternFill
    for c in ws[row]:
        c.font = Font(name="Calibri", bold=True)
        c.fill = PatternFill("solid", fgColor=fill)


def write_l08(path, x):
    from openpyxl import Workbook
    mon = P.monthly(P.study_counts(x, "L08"))
    wb = Workbook()
    ws = wb.active
    ws.title = "2026 by month"
    ws.append(["Lanthorne Field (L08): 2026 operations count"])
    ws.append(["Pellam Airfield Services for the %s. Counted around the clock, local time." % W.DISTRICT])
    ws.append([])
    ws.append(["Month", "Piston airplane takeoffs", "Piston airplane landings", "Piston airplane touch-and-go circuits",
               "Turbine airplane operations", "Helicopter operations", "Glider launches (aero-tow)",
               "Glider landings"])
    _head(ws, 4, fill="E2EFDA")
    for m in range(1, 13):
        v = mon.get(m, {})
        ws.append([MONTHS[m - 1], v.get("takeoffs", 0), v.get("landings", 0), v.get("circuits", 0),
                   v.get("turbine", 0), v.get("helicopter", 0), v.get("launches", 0), v.get("glider_landings", 0)])
    ws.append(["Total"] + ["=SUM(%s5:%s16)" % (c, c) for c in "BCDEFGH"])
    for c in ws[17]:
        c.font = c.font.copy(bold=True)
    for col, w in zip("ABCDEFGH", (12, 14, 14, 18, 15, 13, 15, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B5"
    ws2 = wb.create_sheet("Notes")
    for line in ["How the count was made",
                 "A runway-side counter (radar and camera) logged every runway event; staff reviewed each event's "
                 "images and classed the aircraft.",
                 "A touch-and-go circuit is one pass on the runway without stopping: a landing and a takeoff, two "
                 "operations.",
                 "The tow plane's takeoffs and landings are piston airplane takeoffs and landings; a glider's launch "
                 "and landing are counted separately.",
                 "Prepared 15 January 2027."]:
        ws2.append([line])
    ws2["A1"].font = ws2["A1"].font.copy(bold=True)
    ws2.column_dimensions["A"].width = 110
    wb.save(path)


def write_d19(path, x):
    from openpyxl import Workbook
    daily = P.study_counts(x, "D19")
    wb = Workbook()
    ws = wb.active
    ws.title = "Daily"
    ws.append(["Date", "Piston-engine airplanes", "Turbine airplanes", "Helicopters", "Total"])
    _head(ws, 1, fill="FCE4D6")
    d = dt.date(2026, 1, 1)
    while d.year == 2026:
        v = daily.get(d, {})
        p, t, h = v.get("piston", 0), v.get("turbine", 0), v.get("helicopter", 0)
        ws.append([d, p, t, h, p + t + h])
        d += dt.timedelta(days=1)
    for row in ws.iter_rows(min_row=2, max_col=1):
        for c in row:
            c.number_format = "yyyy-mm-dd"
    for col, w in zip("ABCDE", (12, 22, 17, 12, 9)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("About")
    for line in ["Dunmarsh Airport (D19), 2026 daily operations",
                 "Ardley Counts LLC for the %s." % W.DISTRICT,
                 "Acoustic and video counters at both runway ends, around the clock, 1 January to 31 December 2026; "
                 "every event reviewed. A takeoff or a landing is one operation, a touch-and-go two. Dates are local.",
                 "Delivered 22 January 2027."]:
        ws2.append([line])
    ws2.column_dimensions["A"].width = 110
    wb.save(path)


# ---------------------------------------------------------------------------------------------- state files
def write_based(path, x):
    from openpyxl import Workbook
    sheets = P.based_rows(x)
    wb = Workbook()
    first = True
    for county in W.COUNTIES:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = county
        ws.append(["LocID", "Airport", "N-Number", "Make/Model", "Aircraft Type", "Validated"])
        _head(ws, 1, fill="DDEBF7")
        for r in sheets[county]:
            ws.append(r)
        for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
            for c in row:
                c.number_format = "m/d/yyyy"
        for col, w in zip("ABCDEF", (8, 26, 11, 24, 15, 11)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
    ws = wb.create_sheet("About")
    for line in ["%s: based aircraft validation, 2027" % W.STATE_OFFICE,
                 "Aircraft based at each airport on 1 January 2027, as validated by the airport sponsor against the "
                 "FAA register. One sheet per county.",
                 "Aircraft types follow the national based aircraft inventory: Single Engine, Multi Engine, Jet, "
                 "Helicopter, Glider.",
                 "Released 5 January 2027."]:
        ws.append([line])
    ws.column_dimensions["A"].width = 110
    wb.save(path)


def tower_hours_text(lid):
    o, c, wint = W.TOWER_HOURS[lid]
    s = "%02d%02d-%02d%02d" % (o[0], o[1], c[0], c[1])
    if wint:
        s += " (Nov-Feb %02d%02d-%02d%02d)" % (wint[0][0], wint[0][1], wint[1][0], wint[1][1])
    if lid in W.CONTRACT_TOWERS:
        s += "; closed Dec 25 and Jan 1"
    return s


ELEV = {"HWK": 1248, "CRV": 1105, "OSY": 1390, "VNL": 1512, "KTY": 1176, "DMR": 1324, "F37": 1210, "L08": 1286,
        "G42": 1452, "D19": 1608, "T61": 1150, "P05": 1233, "S27": 1404, "U44": 1677, "O16": 1544, "W33": 1311,
        "Y02": 1720, "A71": 1121, "N26": 1478, "T13": 1590, "K58": 1188, "Z14": 1267, "H52": 1702, "M09": 1395,
        "C66": 1140, "B48": 1325, "D73": 1462, "R31": 1655, "W85": 1102, "L29": 1281}


def directory_rows(x):
    rows = []
    for f in W.F:
        lid = f["lid"]
        rw = P.RUNWAYS[lid]
        fuel = "100LL, Jet A" if (f["towered"] or f["turbines"] >= 3) else ("100LL" if f["B"] >= 60 else "None")
        rows.append([lid, f["name"], f["county"], "Public", "Public" if f["B"] >= 50 or f["towered"] else "Private",
                     "Y" if f["towered"] else "N", tower_hours_text(lid) if f["towered"] else "",
                     "%s/%s" % rw, fuel, "Y" if f["school"] else "N", ELEV[lid]])
    for lid, name, county in W.STRIPS:
        rows.append([lid, name, county, "Private", "Private", "N", "", "", "None", "N", 1100 + 37 * len(name)])
    return rows


def write_directory(path, x):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Airports"
    ws.append(["LocID", "Facility Name", "County", "Use", "Ownership", "ATCT", "ATCT Hours (local)", "Runways", "Fuel",
               "Flight Training", "Elevation (ft)"])
    _head(ws, 1, fill="DDEBF7")
    for r in directory_rows(x):
        ws.append([None if v == "" else v for v in r])
    for col, w in zip("ABCDEFGHIJK", (8, 26, 10, 8, 10, 6, 40, 9, 13, 14, 13)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B2"
    ws2 = wb.create_sheet("Notes")
    for line in ["%s, 2026 Airport Directory (Corrisan Valley airports)" % W.STATE_OFFICE,
                 "ATCT hours are local time. Flight Training: a flight school based on the field.",
                 "Published 1 November 2026."]:
        ws2.append([line])
    ws2.column_dimensions["A"].width = 100
    wb.save(path)


def write_5010(path, x):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "5010 data"
    ws.append(["LocID", "Facility Name", "Based: Single Engine", "Based: Multi Engine", "Based: Jet",
               "Based: Helicopter", "Based: Glider", "Based: Ultralight", "Based: Military", "Ops: Air Taxi",
               "Ops: GA Local", "Ops: GA Itinerant", "Ops: Military", "Operations for 12 months ending",
               "Based aircraft as of"])
    _head(ws, 1, fill="EDEDED")
    for r in P.m5010_rows(x):
        ws.append(r)
    for row in ws.iter_rows(min_row=2, min_col=14, max_col=15):
        for c in row:
            c.number_format = "mm/dd/yyyy"
    for col in "ABCDEFGHIJKLMNO":
        ws.column_dimensions[col].width = 12 if col not in "B" else 26
    ws.freeze_panes = "C2"
    ws2 = wb.create_sheet("Notes")
    for line in ["Airport Master Record (FAA Form 5010) data for the valley's public-use airports, as last filed.",
                 "Based aircraft and annual operations are as each airport reported them at its last 5010 inspection; "
                 "operations are the airport manager's estimate for the 12 months ending on the date shown.",
                 "Copied from the state's airport inspection files, 30 September 2026."]:
        ws2.append([line])
    ws2.column_dimensions["A"].width = 110
    wb.save(path)
