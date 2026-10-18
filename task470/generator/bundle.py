"""Read the shipped pack and build the district's 2026 inventory from it: the network's events in local time, each
event's aircraft through the FAA register, the counted fields' unrecorded operations, the based piston airplanes the
network never recorded, and each field's lead. Only the pack's files are read; the options replay the natural
mishandlings of the pack's mess (checks.py asserts each one leaves the call where it was)."""
import csv
import datetime as dt
import os
import re
import subprocess

import numpy as np
import pandas as pd

import world as W

FIELDS = [f["lid"] for f in W.F]
DST = (pd.Timestamp("2026-03-08 10:00"), pd.Timestamp("2026-11-01 09:00"))


def path(target, key):
    return os.path.join(target, W.FILES[key])


# ---------------------------------------------------------------------------------------------- readers
def events(target, dst=True):
    df = pd.read_csv(path(target, "events"), dtype={"runway": str, "flight_id": str, "icao24": str})
    u = pd.to_datetime(df["event_time_utc"].str.rstrip("Z"))
    off = np.where((u >= DST[0]) & (u < DST[1]), -7, -8) if dst else np.full(len(u), -8)
    df["utc"] = u
    df["local"] = u + pd.to_timedelta(off, unit="h")
    df["date"] = df["local"].dt.date
    df["n_ops"] = np.where(df["event"] == "touch_and_go", 2, 1)
    return df


def padded(p):
    with open(p, encoding="utf-8") as fh:
        rows = list(csv.reader(fh))
    head = [h.strip() for h in rows[0]][:-1]
    return [dict(zip(head, [v.strip() for v in r[:-1]])) for r in rows[1:] if r]


def ymd(s):
    return dt.date(int(s[:4]), int(s[4:6]), int(s[6:8])) if s else None


def register(target):
    master = padded(path(target, "master"))
    dereg = padded(path(target, "dereg"))
    ref = {r["CODE"]: r for r in padded(path(target, "acftref"))}
    eng = {r["CODE"]: r for r in padded(path(target, "engine"))}
    regs = []
    for r in master:
        regs.append(dict(n="N" + r["N-NUMBER"], hex=r["MODE S CODE HEX"].lower(), model=r["MFR MDL CODE"],
                         serial=r["SERIAL NUMBER"], issued=ymd(r["CERT ISSUE DATE"]), cancel=None, current=True,
                         type_acft_reg=r["TYPE AIRCRAFT"]))
    for r in dereg:
        regs.append(dict(n="N" + r["N-NUMBER"], hex=r["MODE S CODE HEX"].lower(), model=r["MFR-MDL-CODE"],
                         serial=r["SERIAL-NUMBER"], issued=ymd(r["CERT-ISSUE-DATE"]), cancel=ymd(r["CANCEL-DATE"]),
                         current=False, type_acft_reg=None))
    for g in regs:
        a = ref[g["model"]]
        g["type_acft"], g["type_eng"], g["engines"] = a["TYPE-ACFT"], a["TYPE-ENG"], int(a["NO-ENG"] or 0)
        g["piston"] = g["type_eng"] == "1" and g["type_acft"] in ("4", "5")
        g["cls"] = ("single" if g["type_acft"] == "4" else "multi") if g["piston"] else "other"
    return dict(regs=regs, ref=ref, eng=eng)


def held(g, d):
    """Was registration g in force on local date d?"""
    return g["issued"] <= d and (g["cancel"] is None or d < g["cancel"])


def resolve(R, how="broadcast", dereg=True):
    """A function (hex, local date, flight ID) -> registration. 'asof': the registration that held the address on that
    date by the register's dates (a cancelled one by its issue and cancel dates); 'broadcast': the same, except that an
    operation whose flight ID is the N-number of another airframe's registration in force that day is that airframe's
    (the aircraft that flew it); 'current': today's register, falling back to the latest cancelled registration for
    an address it no longer lists; dereg=False: today's register only."""
    by_hex, by_n = {}, {}
    for g in R["regs"]:
        by_hex.setdefault(g["hex"], []).append(g)
        by_n.setdefault(g["n"], []).append(g)
    cache = {}

    def f(h, d, fid=None):
        k = (h, d, fid) if how == "broadcast" else ((h, d) if how == "asof" else h)
        if k in cache:
            return cache[k]
        gs = by_hex.get(h, [])
        cur = [g for g in gs if g["current"]]
        old = sorted([g for g in gs if not g["current"]], key=lambda g: g["cancel"]) if dereg else []
        out = None
        if how in ("asof", "broadcast"):
            for g in old:
                if g["issued"] <= d < g["cancel"]:
                    out = g
            if out is None and cur:
                out = cur[0]
            if how == "broadcast" and fid:
                for g in by_n.get(norm_n(fid), []):
                    if held(g, d) and (out is None or (g["serial"], g["model"]) != (out["serial"], out["model"])):
                        out = g
        else:
            out = cur[0] if cur else (old[-1] if old else None)
        cache[k] = out
        return out
    return f


def norm_n(s):
    s = str(s).strip().upper().replace("-", "").replace(" ", "")
    return s if s.startswith("N") else "N" + s


def based(target):
    import openpyxl
    wb = openpyxl.load_workbook(path(target, "based"), read_only=True, data_only=True)
    out = []
    for c in W.COUNTIES:
        for row in wb[c].iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            out.append(dict(lid=row[0], n=norm_n(row[2]), raw=row[2], model=row[3], type=row[4], county=c))
    return out


def directory(target):
    import openpyxl
    wb = openpyxl.load_workbook(path(target, "directory"), read_only=True, data_only=True)
    rows = list(wb["Airports"].iter_rows(values_only=True))
    head = rows[0]
    return {r[0]: dict(zip(head, r)) for r in rows[1:] if r[0]}


def tower_window(text):
    """'0700-2100 (Nov-Feb 0700-1900); closed Dec 25 and Jan 1' -> (open, close, winter or None, closed days)."""
    m = re.match(r"(\d{4})-(\d{4})", text)
    o, c = int(m.group(1)), int(m.group(2))
    w = re.search(r"Nov-Feb (\d{4})-(\d{4})", text)
    closed = [(12, 25), (1, 1)] if "closed Dec 25 and Jan 1" in text else []
    return o, c, (int(w.group(1)), int(w.group(2))) if w else None, closed


def in_tower_hours(local, spec, winter=True):
    o, c, w, closed = spec
    hhmm = local.dt.hour * 100 + local.dt.minute
    win = local.dt.month.isin([11, 12, 1, 2])
    if w and winter:
        lo = np.where(win, w[0], o)
        hi = np.where(win, w[1], c)
    else:
        lo, hi = o, c
    ok = (hhmm >= lo) & (hhmm < hi)
    if winter:
        for mo, dd in closed:
            ok &= ~((local.dt.month == mo) & (local.dt.day == dd))
    return ok


def atct(target):
    df = pd.read_csv(path(target, "atct"))
    df["civil"] = df["Itinerant Air Carrier"] + df["Itinerant Air Taxi"] + df["Itinerant General Aviation"] + \
        df["Local Civil"]
    df["all"] = df["Total Operations"]
    return df


def pdf_layout(p):
    return subprocess.run(["pdftotext", "-layout", p, "-"], capture_output=True, text=True, check=True).stdout


def _num(s):
    return int(s.replace(",", ""))


def study_piston(target):
    """Each count study's 2026 piston-engine airplane operations, as the study reports them."""
    import openpyxl
    out = {}
    txt = pdf_layout(path(target, "f37"))
    m = re.search(r"^\s*2026\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s*$", txt, re.M)
    out["F37"] = _num(m.group(1))
    txt = pdf_layout(path(target, "g42"))
    m = re.search(r"^\s*Year\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s*$", txt, re.M)
    out["G42"] = _num(m.group(2))
    wb = openpyxl.load_workbook(path(target, "l08"), read_only=True)
    rows = list(wb["2026 by month"].iter_rows(min_row=5, max_row=16, values_only=True))
    out["L08"] = sum(r[1] + r[2] + 2 * r[3] for r in rows)
    wb = openpyxl.load_workbook(path(target, "d19"), read_only=True)
    out["D19"] = sum(r[1] for r in wb["Daily"].iter_rows(min_row=2, values_only=True) if r[0] is not None)
    return out


# ---------------------------------------------------------------------------------------------- the inventory
class Pack:
    """year: 'local' (golden), 'utc' (the year read on UTC dates) or 'all' (the export's padding days kept);
    dst: False reads every time at UTC-8; winter: False ignores the towers' winter hours and closed days;
    dereg: False prices only addresses today's register lists; civil: 'all' takes the towers' military columns
    as civil; link: False tests a based airplane's recording by its current address only."""

    def __init__(self, target, year="local", dst=True, winter=True, dereg=True, civil="civil", link=True, _cache=None):
        self.target = target
        self.opts = dict(year=year, dst=dst, winter=winter, dereg=dereg, civil=civil, link=link)
        c = _cache or {}
        self.ev = c.get("ev") if (c.get("ev") is not None and dst) else events(target, dst)
        if year == "local":
            self.ev = self.ev[self.ev["local"].dt.year == 2026]
        elif year == "utc":
            self.ev = self.ev[self.ev["utc"].dt.year == 2026]
        self.R = c.get("R") or register(target)
        self.based = c.get("based") or based(target)
        self.dir = c.get("dir") or directory(target)
        self.atct = c.get("atct") if c.get("atct") is not None else atct(target)
        self.study = c.get("study") or study_piston(target)
        self.spec = {lid: tower_window(self.dir[lid]["ATCT Hours (local)"]) for lid in W.TOWERED}

    def cache(self):
        return dict(ev=events(self.target, True), R=self.R, based=self.based, dir=self.dir, atct=self.atct,
                    study=self.study)

    def priced(self, how="broadcast"):
        """The year's events with their aircraft's class and lead (g)."""
        ev = self.ev.copy()
        f = resolve(self.R, how, self.opts["dereg"])
        cls = []
        for h, d, fid in zip(ev["icao24"].values, ev["date"].values, ev["flight_id"].fillna("").values):
            g = f(h, d, fid)
            cls.append(g["cls"] if g else "unknown")
        ev["cls"] = cls
        L = W.LEAD
        word = {"takeoff": "departure", "landing": "arrival", "touch_and_go": "touch_and_go"}
        g = np.zeros(len(ev))
        for c in ("single", "multi"):
            for e, k in word.items():
                sel = (ev["cls"].values == c) & (ev["event"].values == e)
                g[sel] = L[(c, k)] * (2 if e == "touch_and_go" else 1)
        ev["lead_g"] = g
        ev["piston_ops"] = np.where(np.isin(ev["cls"].values, ["single", "multi"]), ev["n_ops"].values, 0)
        return ev

    def heard_hexes(self):
        return set(self.ev["icao24"].unique())

    def unrecorded_tower(self):
        out = {}
        col = "civil" if self.opts["civil"] == "civil" else "all"
        for lid in W.TOWERED:
            e = self.ev[self.ev["site"] == lid]
            heard = int(e["n_ops"][in_tower_hours(e["local"], self.spec[lid], self.opts["winter"])].sum())
            civil = int(self.atct.loc[self.atct["Facility"] == lid, col].sum())
            out[lid] = dict(civil=civil, heard_in_hours=heard, unrecorded=civil - heard)
        return out

    def addresses(self, n, by_n, by_serial):
        """Every address a based airplane carried: its current registration's, and with link=True those of the
        same airframe's registrations cancelled in 2026 (an N-number changed during the year)."""
        gs = by_n.get(n, [])
        cur = [g for g in gs if g["current"]]
        g = cur[0] if cur else (max(gs, key=lambda z: z["cancel"]) if gs else None)
        if g is None:
            return None, set()
        hs = {g["hex"]}
        if self.opts["link"]:
            for o in by_serial.get((g["serial"], g["model"]), []):
                if o["cancel"] and o["cancel"].year >= 2026:
                    hs.add(o["hex"])
        return g, hs

    def never_heard(self):
        """{field: [based piston airplanes the network did not record in 2026]} for the 30 fields, and the based
        piston airplanes."""
        heard = self.heard_hexes()
        by_n, by_serial = {}, {}
        for g in self.R["regs"]:
            by_n.setdefault(g["n"], []).append(g)
            by_serial.setdefault((g["serial"], g["model"]), []).append(g)
        out = {lid: [] for lid in FIELDS}
        piston_based = {lid: [] for lid in FIELDS}
        hexes = {lid: set() for lid in FIELDS}
        for b in self.based:
            if b["lid"] not in out:
                continue
            g, hs = self.addresses(b["n"], by_n, by_serial)
            if g is None or not g["piston"]:
                continue
            piston_based[b["lid"]].append(b["n"])
            hexes[b["lid"]] |= hs
            if not (hs & heard):
                out[b["lid"]].append(b["n"])
        self._based_hexes = hexes
        return out, piston_based

    def based_activity(self, inv):
        """Each field's piston operations at the field per based piston airplane the network recorded, flown by
        those airplanes."""
        ev = self.priced()
        out = {}
        for lid in FIELDS:
            e = ev[(ev["site"] == lid) & ev["icao24"].isin(self._based_hexes[lid])]
            nh = len(inv["piston_based"][lid]) - inv["nd"][lid]
            out[lid] = float(e["piston_ops"].sum()) / nh if nh else 0.0
        return out

    def inventory(self, how="broadcast", k_method="ratio", k_fields=None):
        ev = self.priced(how)
        heard = ev.groupby("site").agg(lead=("lead_g", "sum"), ops=("piston_ops", "sum"),
                                        all_ops=("n_ops", "sum")).reindex(FIELDS).fillna(0)
        unrec = {}
        tw = self.unrecorded_tower()
        for lid in W.TOWERED:
            unrec[lid] = tw[lid]["unrecorded"]
        for lid in W.STUDIED:
            unrec[lid] = self.study[lid] - int(heard.loc[lid, "ops"])
        nd, pb = self.never_heard()
        NDn = {lid: len(v) for lid, v in nd.items()}
        k = k_of(unrec, NDn, k_method, k_fields)
        lead, ops, est = {}, {}, {}
        for lid in FIELDS:
            if lid in unrec:
                u = unrec[lid]
            else:
                u = k * NDn[lid]
                est[lid] = u
            ops[lid] = heard.loc[lid, "ops"] + u
            lead[lid] = (heard.loc[lid, "lead"] + u * W.LEAD_UNRECORDED) / 1000.0
        return dict(lead=lead, ops=ops, k=k, unrec=unrec, nd=NDn, nd_list=nd, piston_based=pb, heard=heard, est=est,
                    tower=tw)


def k_of(unrec, NDn, method="ratio", fields=None):
    C = fields or W.COUNTED
    u = np.array([unrec[l] for l in C], float)
    n = np.array([NDn[l] for l in C], float)
    if method == "ratio":
        return float(u.sum() / n.sum())
    if method == "mean_of_ratios":
        return float(np.mean(u / n))
    if method == "ols_origin":
        return float((u * n).sum() / (n * n).sum())
    if method == "median":
        return float(np.median(u / n))
    raise ValueError(method)


def qualifying(lead, threshold=W.THRESHOLD_KG):
    return sorted(l for l in W.UNTOWERED if lead[l] >= threshold)
