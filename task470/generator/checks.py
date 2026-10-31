"""task470 checks: the three rules executed on the clean truth (R1 to R4), the ladder and its margins, the household-day
family, the semantic fork table (SF1 to SF23), the separation of the asks' devices from the call, the two tables and
their stops, the Shock and back-test screens, every graded figure's distance to its bin edge, the input gates and
distractors, the realism sweeps and the deliverables' agreement with figures.json."""
import datetime as dt
import json
import os
import re
import subprocess
import zipfile

import numpy as np
import pandas as pd

import methods as M
import world as W

RESULTS = []
GOLD = ["5802", "5803", "5805", "5807", "5809"]
ALL = [m[0] for m in W.MUNICIPALITIES]


def check(label, ok, detail=""):
    RESULTS.append((label, bool(ok), detail))
    print("%s  %s%s" % ("ok  " if ok else "FAIL", label, ("  [%s]" % detail) if (detail != "" and not ok) else ""))
    return ok


def pdf_text(path):
    out = subprocess.run(["pdftotext", "-raw", path, "-"], capture_output=True, text=True, check=True).stdout
    return re.sub(r"\s+", " ", out)


def ooxml_text(path):
    with zipfile.ZipFile(path) as z:
        return " ".join(z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.endswith(".xml"))


def shares(vac_by_nr):
    out = {}
    for nr, name, kind, N, *_r in W.MUNICIPALITIES:
        n, v = vac_by_nr[nr]
        out[nr] = round(100.0 * v / n, 1)
    return out


def designate(sh):
    return sorted(nr for nr, s in sh.items() if s >= W.THRESHOLD)


def signature_of(a):
    """a: dwellings x 14 x 24 hourly kWh in local time."""
    with np.errstate(invalid="ignore"):
        night = np.nanmean(a[:, :, 1:5], axis=2)
        rise = np.fmax(np.nanmean(a[:, :, 6:9], axis=2), np.nanmean(a[:, :, 16:22], axis=2)) - night
    return np.nanmax(np.where(np.isnan(rise), -9, rise), axis=1)


# ---------------------------------------------------------------------------------------------- the clean truth
def truth(w):
    """The world before any file is written: each counted dwelling's own hourly load in Norwegian time (no missing
    hours, no estimates, no clock, no keys), whether anybody lived there, and the persons registered there on
    1 November."""
    dws = [u for u in w.units if u["counted"]]
    a = np.array([w.load[u["id"]] for u in dws]).reshape(len(dws), W.DAYS, 24)
    reg = {}
    for p in w.persons:
        if p["nov1_unit"] is not None:
            reg[p["nov1_unit"]] = reg.get(p["nov1_unit"], 0) + 1
    return dict(dws=dws, a=a, vacant=np.array([bool(u["vacant"]) for u in dws]),
                reg=np.array([reg.get(u["id"], 0) for u in dws]), nr=np.array([u["nr"] for u in dws]))


def truth_call(T, vac):
    by = {}
    for nr in ALL:
        m = T["nr"] == nr
        by[nr] = (int(m.sum()), int((vac & m).sum()))
    return designate(shares(by))


# ---------------------------------------------------------------------------------------------- the run
def run(target, golden, w=None, meta_path=None, fig_path=None):
    del RESULTS[:]
    d = M.load(target)
    res, cnt = M.call(d)
    gold = M.designated(res)
    F = json.load(open(fig_path, encoding="utf-8")) if fig_path else None
    files = sorted(os.listdir(target))
    fmts = {f.rsplit(".", 1)[1] for f in files}
    nh = len(d["hourly"])

    # ------------------------------------------------------------------ input gates
    check("gates: %d files (10 or more), formats %s (3 or more), the hourly values %d rows (25,000 or more)" % (
        len(files), sorted(fmts), nh), len(files) >= 10 and len(fmts) >= 3 and nh >= 25000)
    check("gates: the pack is exactly the files world.FILES declares", set(files) == set(W.FILES.values()),
          sorted(set(files) ^ set(W.FILES.values())))
    check("gates: two distractors, the day-ahead prices and the outage log, named in metadata.json",
          meta_path is None or sorted(json.load(open(meta_path))["distractor_files"]) ==
          sorted(W.FILES[k] for k in W.DISTRACTORS))

    # ------------------------------------------------------------------ the call
    check("call: the golden designates %s" % " ".join(GOLD), gold == GOLD, gold)
    raw = {nr: r["raw"] for nr, r in res.items()}
    near = min(abs(v - 4.95) for v in raw.values())
    check("call: no municipality's unrounded share within 0.2 points of the rounding line at 4.95 (closest %.3f)" % near,
          near >= 0.2)
    edge = min(abs((v * 10) % 1 - 0.5) for v in raw.values())
    check("call: no share on a rounding edge (closest %.3f of a tenth)" % edge, edge >= 0.03)
    sig = M.signature(M.series(d)).reindex(cnt["bruksenhet_id"]).fillna(-9).to_numpy()
    vac = cnt["vacant"].to_numpy()
    check("call: every counted dwelling has a metering point with values in the fortnight",
          len(set(M.meters_of(d).values())) == len(cnt), len(cnt) - len(set(M.meters_of(d).values())))
    gap = (sig[~vac].min(), sig[vac].max())
    check("call: the household day separates the dwellings, lived-in minimum %.3f against vacant maximum %.3f" % gap,
          gap[0] - gap[1] > 0.5)

    # ------------------------------------------------------------------ R1 to R4: the three rules
    heads = []
    for f in files:
        if f.endswith(".csv"):
            heads.append(open(os.path.join(target, f), encoding="utf-8").readline().lower())
    heads.append(" ".join(d["hourly"].columns).lower())
    check("R1 the field's raw material ships: the hourly values per metering point, the meter register, the cadastre "
          "and the registers, and no shipped column holds occupancy, vacancy or a household",
          nh > 1000000 and not any(re.search(r"ledig|vacan|occup|bebodd|husstand|tom_", h) for h in heads))
    if w is not None:
        T = truth(w)
        lv = T["a"].sum(axis=(1, 2)) / W.DAYS
        fl = {f: truth_call(T, lv < f) for f in range(1, 11)}
        check("R2 analyst test on the clean truth: the utility-data vacancy floor, any floor from 1 to 10 kWh a day on "
              "each dwelling's own load, designates no municipality (the default path's Ends on)",
              all(v == [] for v in fl.values()), fl)
        st = signature_of(T["a"])
        check("R3 the decisive step on the clean truth: the household day on any day of the fortnight finds exactly "
              "the dwellings nobody lived in and designates the golden five",
              ((st < M.CUT) == T["vacant"]).all() and truth_call(T, st < M.CUT) == GOLD)
        cheap = {"the register's empty dwellings": truth_call(T, T["reg"] == 0),
                 "the register's empty dwellings with no household day": truth_call(T, (T["reg"] == 0) & (st < M.CUT)),
                 "the register's empty dwellings under 3 kWh a day": truth_call(T, (T["reg"] == 0) & (lv < 3))}
        for f in (1, 3, 5, 10, 15, 20):
            cheap["a floor of %d kWh a day" % f] = truth_call(T, lv < f)
        landing = [k for k, v in cheap.items() if v == GOLD]
        check("R4 Gate F on the clean truth: of %d cheap routes none lands the call; the household day is the one "
              "route that does" % len(cheap), not landing, landing)
        check("R4 the register designates all ten and the register's empty dwellings without a household day only "
              "Hesteby, on the clean truth", cheap["the register's empty dwellings"] == ALL and
              cheap["the register's empty dwellings with no household day"] == ["5803"])

    # ------------------------------------------------------------------ the ladder on the pack
    reg, _ = M.call(d, "register")
    check("ladder rung 0: the population register's empty dwellings on 1 November designate all ten",
          M.designated(reg) == ALL, M.designated(reg))
    sets = {}
    for f in np.arange(0.5, 30.01, 0.5):
        r, _ = M.call(d, "level", float(f))
        sets[float(f)] = M.designated(r)
    check("ladder rung 1: a consumption floor from 0.5 to 10 kWh a day designates none",
          all(v == [] for f, v in sets.items() if f <= 10.0))
    check("ladder rung 1: no floor from 0.5 to 30 kWh a day designates the golden set", all(v != GOLD for v in
                                                                                           sets.values()))
    check("ladder rung 1: floors of 14.5 kWh a day and more designate all ten",
          all(v == ALL for f, v in sets.items() if f >= 14.5))
    for f in (2, 5, 10):
        r, _ = M.call(d, "reg_and_level", float(f))
        check("ladder: the register's empty dwellings under %d kWh a day designate none" % f, M.designated(r) == [])
    rs, _ = M.call(d, "reg_and_shape")
    check("ladder rung 2: the register's empty dwellings with no household day designate Hesteby alone",
          M.designated(rs) == ["5803"], M.designated(rs))

    # ------------------------------------------------------------------ the household-day family (SF12, SF13)
    g = M.series(d).reindex(cnt["bruksenhet_id"])
    a = g.to_numpy().reshape(len(g), W.DAYS, 24)
    with np.errstate(invalid="ignore"):
        night = np.nanmean(a[:, :, 1:5], axis=2)
        morn = np.nanmean(a[:, :, 6:9], axis=2) - night
        eve = np.nanmean(a[:, :, 16:22], axis=2) - night
        rng_ = np.nanmax(a, axis=2) - np.nanmin(a, axis=2)
        sd = np.nanstd(a, axis=2)
    fam = {"morning or evening over night, any day": np.nanmax(np.fmax(morn, eve), axis=1),
           "evening over night, any day": np.nanmax(eve, axis=1),
           "morning over night, any day": np.nanmax(morn, axis=1),
           "the day's highest hour over its lowest, any day": np.nanmax(rng_, axis=1),
           "the day's spread of hours, any day": np.nanmax(sd, axis=1)}
    for k, v in fam.items():
        lo, hi = np.nanmin(v[~vac]), np.nanmax(v[vac])
        if "highest" in k or "spread" in k:
            check("SF12 the household day as %s with the outage hours read as use: lived-in minimum %.3f, vacant "
                  "maximum %.3f, the outage log's hours with no supply opening the vacant days" % (k, lo, hi), hi > lo)
            continue
        check("SF12 the household day as %s: lived-in minimum %.3f above vacant maximum %.3f, so every cut between "
              "gives the golden dwellings" % (k, lo, hi), lo > hi)
    out = _outage_mask(target, cnt)
    b = np.where(out, np.nan, a)
    with np.errstate(invalid="ignore"):
        for k, v in (("the day's highest hour over its lowest", np.nanmax(np.nanmax(b, axis=2) - np.nanmin(b, axis=2),
                                                                        axis=1)),
                     ("the day's spread of hours", np.nanmax(np.nanstd(b, axis=2), axis=1))):
            lo, hi = np.nanmin(v[~vac]), np.nanmax(v[vac])
            check("SF12 the household day as %s, the outage log's hours left out: lived-in minimum %.3f above vacant "
                  "maximum %.3f" % (k, lo, hi), lo > hi)
    for cut in (0.2, 0.3, 0.5, 0.7, 0.8):
        r = cnt.assign(vacant=(fam["morning or evening over night, any day"] < cut))
        check("SF12 a cut of %.1f kWh designates the golden five" % cut,
              designate({nr: round(100.0 * r[r.kommunenr == nr].vacant.sum() / (r.kommunenr == nr).sum(), 1)
                         for nr in ALL}) == GOLD)
    mean_rise = np.nanmean(np.fmax(morn, eve), axis=1)
    mv = cnt.assign(vacant=mean_rise < M.CUT)
    msets = designate({nr: round(100.0 * mv[mv.kommunenr == nr].vacant.sum() / (mv.kommunenr == nr).sum(), 1)
                       for nr in ALL})
    check("SF13 the household day averaged over the fortnight (ruled out by Section 4, on any day) designates %s" %
          " ".join(msets), True)

    # ------------------------------------------------------------------ the semantic fork table
    cad = d["cad"]
    blank = cad[(cad.bruksenhetstype == "Bolig") & (cad.tatt_i_bruk == "")]
    late = cad[(cad.bruksenhetstype == "Bolig") & (cad.tatt_i_bruk != "") & (M._date(cad.tatt_i_bruk) > M.REF)]
    check("SF1 Section 3 counts Bolig units in use on or before 1 November: %d units with no date and %d taken into "
          "use later are left out, and holiday homes and care homes have their own type" % (len(blank), len(late)),
          len(blank) > 0 and len(late) > 0 and len(cnt) == sum(m[3] for m in W.MUNICIPALITIES))
    check("SF2 Section 4: the register's reading designates all ten, so the sentence that the register does not decide "
          "is what rules it out", M.designated(reg) != GOLD)
    utc, _ = _call_with(d, clock="utc")
    check("SF3 the fortnight on UTC instead of Norwegian time: same call", M.designated(utc) == GOLD)
    for how in ("first", "both"):
        r, _ = _call_with(d, dedupe=how)
        check("SF4 an estimated hour and its measured replacement, %s kept: same call" % how, M.designated(r) == GOLD)
    r, _ = _call_with(d, meters="current")
    check("SF5 a dwelling whose meter changed in the fortnight read from its current point only: same call",
          M.designated(r) == GOLD)
    fp = d["mp"][d["mp"].kundekategori == "Fellesmåling"]
    byadr = {}
    for x in cnt.itertuples():
        byadr.setdefault((x.kommunenr, M.norm_addr("%s %s%s" % (x.adressenavn, x.husnr, x.bokstav))), []).append(x)
    uniq = sum(1 for x in fp.itertuples() if len(byadr.get((x.kommunenr, M.norm_addr(x.adresse)), [])) == 1)
    r, _ = _call_with(d, felles=True)
    check("SF6 the communal points (Fellesmåling) joined by address wherever one dwelling has it (%d of %d, under "
          "Snøholt's old number): same call" % (uniq, len(fp)), len(fp) > 0 and M.designated(r) == GOLD)
    sn = d["mp"][(d["mp"].kundekategori == "Husholdning") & (d["mp"].kommunenr == "5809")]
    sn_ids = set(cnt[cnt.kommunenr == "5810"].bruksenhet_id)
    check("SF7 Snøholt's points carry Fjordeid's old number 5809, and every Snøholt dwelling's point carries its unit "
          "ID, so the join by ID places them", (d["mp"].kommunenr == "5810").sum() == 0 and
          sn_ids <= set(sn.bruksenhet_id))
    unr = designate({nr: res[nr]["raw"] for nr in ALL})
    check("SF8 the threshold on the unrounded share instead of the rounded: same call", unr == GOLD)
    check("SF9 the ten municipalities as they stood on 1 January 2026, Snøholt split from Fjordeid in 2024",
          sorted(d["mun"].kommunenr) == ALL)
    r, _ = _call_with(d, gaps="skip_day")
    check("SF10 a day with a missing hour left out of the household day instead of read on the hours present: same "
          "call", M.designated(r) == GOLD)
    r, _ = _call_with(d, outages="drop")
    check("SF11 the outage hours dropped instead of read as zero: same call", M.designated(r) == GOLD)
    two = cad[(cad.bygningstype == "Tomannsbolig")]
    check("SF14 a two-unit house is two dwellings with a metering point each", len(two) > 0)
    check("SF15 no holiday home or care home has a household metering point in the extract",
          not set(d["mp"].bruksenhet_id) & set(cad[cad.bruksenhetstype != "Bolig"].bruksenhet_id))

    # ------------------------------------------------------------------ separation of the asks from the call
    call_only = {k: d[k] for k in ("cad", "mp", "hourly", "mun")}
    r, _ = M.call(call_only)
    check("separation: the call computed from the cadastre, the meter register and the hourly values alone",
          M.designated(r) == GOLD)
    ren_idx, relabel = M.renumberings(d)
    newk = {v[0] for v in relabel.values()}
    ground = {k for k in relabel if k[2].startswith("H01")}
    cadk = {M._key(x.kommunenr, x.adressenavn, x.husnr, x.bokstav, x.bruksenhetsnr) for x in cad.itertuples()}
    check("separation: the cadastre carries the renumbered blocks' new unit numbers only, no flat on the old ground-floor "
          "numbers (%d relabelled units, %d of them on the old ground floor)" % (len(relabel), len(ground)),
          newk <= cadk and not (ground & cadk) and len(ground) > 0)
    mp_ = d["mp"]
    mpk = {(x.kommunenr, M.norm_addr(x.adresse), M.norm_unit(x.bruksenhetsnr)) for x in mp_.itertuples()}
    mpk |= {("5810", a, u) for nr_, a, u in mpk if nr_ == "5809"}
    check("separation: the meter register carries no old ground-floor number of a renumbered block, so the call's files "
          "hold none of table A's device rows", not (ground & mpk))

    # ------------------------------------------------------------------ table A
    if w is not None:
        ta = {nr: {"u18": 0, "18-66": 0, "67+": 0} for nr in ALL}
        for p in w.persons:
            if p["nov1_unit"] is None or not w.by_id[p["nov1_unit"]]["counted"]:
                continue
            age = M._age(pd.Timestamp(p["born"]), M.REF)
            ta[w.by_id[p["nov1_unit"]]["nr"]]["u18" if age <= 17 else ("18-66" if age <= 66 else "67+")] += 1
        A = M.ask_a(d)
        check("table A: the golden equals the clean truth in all 30 cells", A == ta)
        stops = {"the register as it stands on 5 January": M.ask_a(d, "snapshot"),
                 "every event undone as written, the renumberings as moves (the cracker)": M.ask_a(d, "book"),
                 "moves undone, the dead and the emigrated not added back": M.ask_a(d, "no_deaths"),
                 "ages on 5 January": M.ask_a(d, "golden", age_on="snap")}
        for k, v in stops.items():
            hits = sum(v[nr][b] == ta[nr][b] for nr in ALL for b in ta[nr])
            want = 0 if "cracker" in k else 15
            check("table A stop, %s: %d of 30 cells" % (k, hits), hits <= want)
        ev = d["ev"]
        straddle = ev[(M._date(ev.hendelsesdato) <= M.REF) & (M._date(ev.registrert_dato) > M.REF)]
        check("SF16 no event dated on or before 1 November was registered after it, nor the reverse: the register's "
              "date and the event's date select the same persons", len(straddle) == 0 and
              ((M._date(ev.registrert_dato) >= M._date(ev.hendelsesdato)).all()))
        check("SF17 the ministry's letter asks for age on 1 November; ages on 5 January miss %d of 30 cells" % (
            30 - sum(stops["ages on 5 January"][nr][b] == ta[nr][b] for nr in ALL for b in ta[nr])),
            stops["ages on 5 January"] != ta)
        check("SF18 the block renumberings: on one day every resident of the block moves to the unit number one floor "
              "up at the same address, and the cadastre holds only the new numbers (%d events)" % len(ren_idx),
              all(_floor_up(k, v[0]) for k, v in relabel.items()))

    # ------------------------------------------------------------------ table B
    if w is not None:
        tb = {}
        for nr in ALL:
            hs = set(str(x) for x in w.muni_holders[nr]) | set(w.muni_legal[nr])
            tb[nr] = sum(1 for u in w.units if u["nr"] == nr and u["counted"] and
                         w.title_nov1.get((u["nr"], u["gnr"], u["bnr"], u["snr"])) in hs)
        B = M.ask_b(d)
        check("table B: the golden equals the clean truth in all ten", B == tb, (B, tb))
        stops = {"organisation numbers only, the KF and the full chain, registered dates (the cracker)":
                 M.ask_b(d, names="none"),
                 "titles without a number matched to the business register by exact name (the cracker)":
                 M.ask_b(d, names="exact"),
                 "every holder whose name begins with the municipality's": M.ask_b(d, names="startswith"),
                 "the titles as they stand on 5 January": M.ask_b(d, basis="snapshot"),
                 "transfers dated by the deed": M.ask_b(d, basis="dokument"),
                 "a single holder with all the shares": M.ask_b(d, chain="single"),
                 "the municipality and its KF only": M.ask_b(d, chain="none"),
                 "no municipal enterprise": M.ask_b(d, kf=False),
                 "titles instead of dwellings": M.ask_b(d, unit="title")}
        for k, v in stops.items():
            hits = sum(v[nr] == tb[nr] for nr in ALL)
            want = 0 if "cracker" in k else 7
            check("table B stop, %s: %d of 10" % (k, hits), hits <= want)
        trf = d["trf"]
        check("SF19 no transfer was registered on 1 November itself", (M._date(trf.tinglyst_dato) != M.REF).all())
        check("SF20 a company the municipality owns with its holding company between them counts (%d such)" % sum(
            1 for nr in ALL if M.municipal_holders(d)[nr] != M.municipal_holders(d, chain="single")[nr]),
            any(M.municipal_holders(d)[nr] != M.municipal_holders(d, chain="single")[nr] for nr in ALL))
        tit = d["tit"]
        noid = tit[(tit.hjemmelshaver_type == "Organisasjon") & (tit.hjemmelshaver_id == "")]
        check("SF21 every title without an organisation number is held in a municipality's own name (%d titles), and "
              "no such name is in the business register" % len(noid),
              len(noid) > 0 and all(any(M.own_name(x.hjemmelshaver_navn, m[1]) for m in W.MUNICIPALITIES)
                                    for x in noid.itertuples())
              and not set(noid.hjemmelshaver_navn) & set(d["ent"].navn))
        check("SF22 the table counts dwellings: titles instead of dwellings miss all ten",
              sum(stops["titles instead of dwellings"][nr] == tb[nr] for nr in ALL) == 0)
        iks = d["ent"][d["ent"].organisasjonsform == "IKS"]
        check("SF23 the inter-municipal company holds dwellings and is no company with all its shares held by one "
              "municipality", len(iks) == 1 and (tit.hjemmelshaver_id == iks.organisasjonsnummer.iloc[0]).sum() > 0)

    # ------------------------------------------------------------------ Shock and back-test
    stat = {}
    reg_n = M.registered_nov1(d).reindex(cnt["bruksenhet_id"]).fillna(0).to_numpy()
    lvp = M.level(M.series(d)).reindex(cnt["bruksenhet_id"]).fillna(0).to_numpy()
    for nr in ALL:
        m = (cnt.kommunenr == nr).to_numpy()
        stat.setdefault("register-empty share", {})[nr] = (reg_n[m] == 0).mean()
        stat.setdefault("mean daily use", {})[nr] = lvp[m].mean()
        stat.setdefault("median daily use", {})[nr] = np.median(lvp[m])
        stat.setdefault("share of flats", {})[nr] = (cad[(cad.kommunenr == nr) & (cad.bruksenhetstype == "Bolig")]
                                                    .bygningstype == "Boligblokk").mean()
        stat.setdefault("counted dwellings", {})[nr] = m.sum()
        for f in (1, 3, 5, 10, 15, 20, 30):
            stat.setdefault("share under %d kWh a day" % f, {})[nr] = (lvp[m] < f).mean()
            stat.setdefault("share over %d kWh a day" % f, {})[nr] = (lvp[m] > f).mean()
        stat.setdefault("holiday homes per dwelling", {})[nr] = ((cad.kommunenr == nr) & (
            cad.bruksenhetstype == "Fritidsbolig")).sum() / m.sum()
    interval = [k for k, v in stat.items() for top in (True, False)
                if sorted(sorted(v, key=lambda nr: v[nr], reverse=top)[:5]) == GOLD]
    check("Shock: the golden five are the top or bottom five of none of %d shipped statistics" % len(stat), not interval,
          interval)
    low = lvp < 3
    bt = ((reg_n[low] == 0).mean(), (reg_n[~low] == 0).mean())
    check("Back-test: the dwellings a 3 kWh floor calls vacant are %.0f per cent register-empty against %.0f per cent of "
          "the rest, so the floor validates against the register" % (100 * bt[0], 100 * bt[1]), bt[0] > 5 * bt[1])

    # ------------------------------------------------------------------ deliverables and figures.json
    if F is not None and golden:
        xl = os.path.join(golden, W.GOLDEN["xlsx"])
        dx = os.path.join(golden, W.GOLDEN["docx"])
        check("deliverables: the workbook and the letter exist", os.path.exists(xl) and os.path.exists(dx))
        from openpyxl import load_workbook
        wb = load_workbook(xl, read_only=True)
        rows = list(wb["Dwellings"].iter_rows(min_row=2, values_only=True))
        check("deliverables: the workbook lists the %s counted dwellings, %s vacant, %s persons" % (
            F["totals"]["counted"], F["totals"]["vacant"], F["totals"]["persons"]),
            len(rows) == len(cnt) and sum(r[6] == "Yes" for r in rows) == int(F["totals"]["vacant"].replace(",", ""))
            and sum(r[5] for r in rows) == int(F["totals"]["persons"].replace(",", "")))
        summ = {str(r[0]): r for r in wb["By municipality"].iter_rows(min_row=2, values_only=True) if r[0]}
        check("deliverables: the workbook's summary agrees with figures.json", all(
            summ[m["nr"]][3] == m["vacant"] and "%.1f" % summ[m["nr"]][4] == m["share"] for m in F["municipalities"]))
        tx = ooxml_text(dx)
        need = [m["share"] for m in F["municipalities"]] + [r["held"] for r in F["ask_b"]] + \
            [r["under_18"] for r in F["ask_a"]] + [r["name"] for r in F["qualifying"]]
        check("deliverables: the letter carries every share, both tables and the qualifying names",
              all(x in tx for x in need))
        check("deliverables: figures.json's per-dwelling persons total equals table A's", int(
            F["totals"]["persons"].replace(",", "")) == sum(int(r[k].replace(",", "")) for r in F["ask_a"]
                                                           for k in ("under_18", "age_18_66", "age_67")))

    # ------------------------------------------------------------------ realism
    bad = []
    for f in files:
        b = open(os.path.join(target, f), "rb").read()
        if "—".encode() in b or "–".encode() in b:
            bad.append(f)
        if f.endswith(".csv") and b"\r\n" not in b[:2000]:
            bad.append(f + " line ends")
    check("realism: no dash characters in any file, every CSV in the registers' CRLF layout", not bad, bad)
    reg_txt = pdf_text(os.path.join(target, W.FILES["regulation"]))
    eml = open(os.path.join(target, W.FILES["assignment"]), encoding="utf-8").read()
    words = r"\b(hour of the day|morning|evening|night|daily|profile|pattern|floor|threshold of|kwh|routine|heated)\b"
    check("realism: neither governing document names an hour of the day, a profile, a consumption floor or heating",
          not re.search(words, (reg_txt + " " + eml).lower()), re.findall(words, (reg_txt + " " + eml).lower()))
    late_dates = []
    for f in files:
        if f.endswith(".csv"):
            for s in re.findall(rb"\b\d{2}\.\d{2}\.(\d{4})\b", open(os.path.join(target, f), "rb").read()[:3000000]):
                if int(s) > 2027:
                    late_dates.append(f)
                    break
    check("realism: no file carries a date after the certification", not late_dates, late_dates)
    ok = all(r[1] for r in RESULTS)
    print("%d checks, %d failed" % (len(RESULTS), sum(not r[1] for r in RESULTS)))
    return ok


def _outage_mask(target, cnt):
    """The hours the grid company's outage log gives with no supply, for every counted dwelling of the municipality
    named: dwellings x 14 x 24, True where the hour falls in an outage."""
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(target, W.FILES["outages"]), read_only=True)
    mask = np.zeros((len(cnt), W.DAYS, 24), bool)
    nrs = cnt["kommunenr"].to_numpy()
    for r in list(wb.worksheets[0].iter_rows(min_row=2, values_only=True)):
        day = (dt.datetime.strptime(r[1], "%d.%m.%Y").date() - W.FORTNIGHT_START).days
        h0, h1 = int(r[2][:2]), int(r[3][:2])
        mask[nrs == str(r[0]), day, h0:h1] = True
    return mask


def _floor_up(old, new):
    return old[:2] == new[:2] and int(new[2][1:3]) == int(old[2][1:3]) + 1 and new[2][3:] == old[2][3:]


def _call_with(d, clock="local", dedupe="latest", meters="all", gaps="nan", outages="zero", felles=False):
    """The call under a variant reading of one convention."""
    h = d["hourly"]
    if dedupe == "both":
        hh = h
    else:
        hh = h.sort_values(["malepunkt_id", "fra_tidspunkt_utc", "leveranse"]).drop_duplicates(
            ["malepunkt_id", "fra_tidspunkt_utc"], keep="last" if dedupe == "latest" else "first")
    m2d = dict(M.meters_of(d))
    if felles:
        byadr = {}
        for x in M.counted(d).itertuples():
            byadr.setdefault((x.kommunenr, M.norm_addr("%s %s%s" % (x.adressenavn, x.husnr, x.bokstav))), []).append(
                x.bruksenhet_id)
        for x in d["mp"][d["mp"].kundekategori == "Fellesmåling"].itertuples():
            c = byadr.get((x.kommunenr, M.norm_addr(x.adresse)), [])
            if len(c) == 1:
                m2d[int(x.malepunkt_id)] = c[0]
    if meters == "current":
        cur = set(d["mp"][d["mp"].gyldig_til == ""].malepunkt_id.astype(np.int64))
        m2d = {k: v for k, v in m2d.items() if k in cur}
    hh = hh[hh["malepunkt_id"].isin(list(m2d))].copy()
    hh["dw"] = hh["malepunkt_id"].map(m2d)
    t = hh["fra_tidspunkt_utc"].dt.tz_convert(None)
    if clock == "local":
        t = t + pd.Timedelta(hours=W.UTC_OFFSET)
    hh["hix"] = ((t - pd.Timestamp(W.FORTNIGHT_START)) / pd.Timedelta(hours=1)).astype(int)
    hh = hh[(hh["hix"] >= 0) & (hh["hix"] < W.HOURS)]
    g = hh.groupby(["dw", "hix"])["forbruk_kwh"].sum().unstack("hix").reindex(columns=range(W.HOURS))
    a = g.to_numpy().reshape(len(g), W.DAYS, 24)
    if gaps == "skip_day":
        bad = np.isnan(a).any(axis=2)
        a = np.where(bad[:, :, None], np.nan, a)
    if outages == "drop":
        a = np.where(a == 0.0, np.nan, a)
    s = pd.Series(signature_of(a), index=g.index)
    cnt = M.counted(d)
    v = s.reindex(cnt["bruksenhet_id"]).fillna(-9).to_numpy() < M.CUT
    out = {}
    for nr in ALL:
        m = (cnt.kommunenr == nr).to_numpy()
        out[nr] = {"share": round(100.0 * v[m].sum() / m.sum(), 1)}
    return out, cnt
