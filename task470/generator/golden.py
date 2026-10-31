"""The golden deliverables, computed from the shipped files only: county_dwelling_occupancy_2026-11.xlsx (every dwelling
the regulation counts, its municipality, the persons registered there on 1 November 2026 and whether it stood vacant,
with the municipalities' summary) and vacancy_grant_certification_2027.docx (the office's certification letter to the
ministry: the qualifying municipalities, each municipality's vacant dwellings and share, the chart against the
threshold and the two tables the ministry asked for); and figures.json, every figure the write-up states."""
import io
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import numpy as np      # noqa: E402

import methods as M     # noqa: E402
import world as W       # noqa: E402

BLUE = "#1f4e79"
GREY = "#a6adb4"
RED = "#9c2f2f"
WORDS = {0: "No", 1: "One", 2: "Two", 3: "Three", 4: "Four", 5: "Five", 6: "Six", 7: "Seven", 8: "Eight", 9: "Nine",
         10: "Ten"}


def fmt(n):
    return "{:,}".format(int(n))


def compute(target):
    d = M.load(target)
    res, cnt = M.call(d)
    per = M.persons_nov1(d).groupby("dw").size()
    cnt = cnt.assign(persons=cnt["bruksenhet_id"].map(per).fillna(0).astype(int))
    g = M.series(d)
    lv = M.level(g).reindex(cnt["bruksenhet_id"]).fillna(0.0).to_numpy()
    floors = {}
    for fl in range(1, 11):
        r, _ = M.call(d, "level", float(fl))
        floors[fl] = M.designated(r)
    reg, _ = M.call(d, "register")
    rs, _ = M.call(d, "reg_and_shape")
    regn = cnt["persons"].to_numpy()
    vac = cnt["vacant"].to_numpy()
    return dict(d=d, res=res, cnt=cnt, lv=lv, floors=floors, reg=reg, rs=rs, ask_a=M.ask_a(d), ask_b=M.ask_b(d),
                vac_reg=int(((regn > 0) & vac).sum()), lived_noreg=int(((regn == 0) & ~vac).sum()),
                renum=len(M.renumberings(d)[1]))


def figures(X):
    res = X["res"]
    names = {nr: res[nr]["name"] for nr in res}
    q = M.designated(res)
    out = [nr for nr in res if nr not in q]
    below = max(out, key=lambda nr: res[nr]["raw"])
    above = min(q, key=lambda nr: res[nr]["raw"])
    lv, vac = X["lv"], X["cnt"]["vacant"].to_numpy()
    assert all(v == [] for v in X["floors"].values())
    return dict(
        count=len(q),
        count_word=WORDS[len(q)],
        threshold="5.0",
        qualifying=[dict(nr=nr, name=names[nr], vacant=res[nr]["vacant"], share="%.1f" % res[nr]["share"]) for nr in q],
        municipalities=[dict(nr=nr, name=names[nr], counted=fmt(res[nr]["counted"]), vacant=res[nr]["vacant"],
                             share="%.1f" % res[nr]["share"], qualifies=nr in q) for nr in res],
        lowest_qualifying=dict(name=names[above], share="%.1f" % res[above]["share"]),
        highest_below=dict(name=names[below], share="%.1f" % res[below]["share"]),
        totals=dict(counted=fmt(sum(r["counted"] for r in res.values())), vacant=fmt(sum(r["vacant"] for r in
                    res.values())), persons=fmt(int(X["cnt"]["persons"].sum()))),
        vacant_registered=fmt(X["vac_reg"]),
        lived_unregistered=fmt(X["lived_noreg"]),
        vacant_median_kwh="%.1f" % np.median(lv[vac]),
        lived_median_kwh="%.1f" % np.median(lv[~vac]),
        rivals=dict(register_count=len(M.designated(X["reg"])), floor_low="1", floor_high="10",
                    register_and_profile=[names[nr] for nr in M.designated(X["rs"])]),
        cut="0.5",
        age_bands=dict(under="18", upper="66", over="67"),
        renumbered=X["renum"],
        ask_a=[dict(nr=nr, name=names[nr], under_18=fmt(a["u18"]), age_18_66=fmt(a["18-66"]), age_67=fmt(a["67+"]))
               for nr, a in X["ask_a"].items()],
        ask_b=[dict(nr=nr, name=names[nr], held=fmt(v)) for nr, v in X["ask_b"].items()],
    )


# ---------------------------------------------------------------------------------------------- the workbook
def write_xlsx(X, F, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Dwellings"
    head = ["Municipality no.", "Municipality", "Dwelling ID", "Address", "Unit", "Persons registered 1 Nov 2026",
            "Vacant 9-22 Nov 2026"]
    ws.append(head)
    names = {m["nr"]: m["name"] for m in F["municipalities"]}
    c = X["cnt"].copy()
    c["addr"] = [("%s %s%s" % (a, n, b)).strip() for a, n, b in zip(c["adressenavn"], c["husnr"], c["bokstav"])]
    c["husnr_i"] = c["husnr"].astype(int)
    c = c.sort_values(["kommunenr", "adressenavn", "husnr_i", "bokstav", "bruksenhetsnr", "bruksenhet_id"])
    for r in c.itertuples():
        ws.append([int(r.kommunenr), names[r.kommunenr], int(r.bruksenhet_id), r.addr, r.bruksenhetsnr, int(r.persons),
                   "Yes" if r.vacant else "No"])
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="DCE6F1")
    for cell in ws[1]:
        cell.font = bold
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col, wd in zip("ABCDEFG", (14, 14, 13, 26, 8, 15, 13)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:G%d" % (len(c) + 1)
    s = wb.create_sheet("By municipality")
    s.append(["Municipality no.", "Municipality", "Counted dwellings", "Vacant dwellings", "Vacancy share (%)",
              "Qualifies"])
    for m in F["municipalities"]:
        s.append([int(m["nr"]), m["name"], int(m["counted"].replace(",", "")), m["vacant"], float(m["share"]),
                  "Yes" if m["qualifies"] else "No"])
    s.append([None, "County", int(F["totals"]["counted"].replace(",", "")), int(F["totals"]["vacant"].replace(",", "")),
              None, None])
    for cell in s[1]:
        cell.font = bold
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in s[s.max_row]:
        cell.font = bold
    for row in s.iter_rows(min_row=2, min_col=5, max_col=5):
        for cell in row:
            cell.number_format = "0.0"
    for col, wd in zip("ABCDEF", (14, 14, 11, 10, 11, 10)):
        s.column_dimensions[col].width = wd
    s.freeze_panes = "A2"
    wb.save(path)


# ---------------------------------------------------------------------------------------------- the letter
def chart_png(F):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    ms = sorted(F["municipalities"], key=lambda m: float(m["share"]))
    vals = [float(m["share"]) for m in ms]
    fig, ax = plt.subplots(figsize=(6.4, 3.6), dpi=150)
    cols = [BLUE if m["qualifies"] else GREY for m in ms]
    ax.barh(range(len(ms)), vals, color=cols, height=0.62)
    ax.axvline(5.0, color=RED, linewidth=1.0, linestyle=(0, (4, 2)))
    ax.text(5.07, len(ms) - 0.45, "Threshold 5.0 %", color=RED, fontsize=7.4, va="center")
    for i, v in enumerate(vals):
        ax.text(v + 0.08, i, "%.1f" % v, va="center", fontsize=7.2, color="#2b2b2b")
    ax.set_yticks(range(len(ms)))
    ax.set_yticklabels(["%s (%s)" % (m["name"], m["nr"]) for m in ms], fontsize=7.6)
    ax.set_xlabel("Vacant dwellings, per cent of counted dwellings, 9 to 22 November 2026", fontsize=7.6)
    ax.set_xlim(0, max(vals) * 1.15)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(axis="x", labelsize=7.4)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def _page_field(paragraph):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _join(xs):
    return ", ".join(xs[:-1]) + " and " + xs[-1] if len(xs) > 1 else xs[0]


def _table(doc, head, rows, right_from=2, widths=None):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Light List Accent 1"
    for c, txt in zip(t.rows[0].cells, head):
        c.text = txt
    for r in rows:
        cells = t.add_row().cells
        for c, v in zip(cells, r):
            c.text = str(v)
    for rw in t.rows:
        for k, c in enumerate(rw.cells):
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                if k >= right_from:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return t


def write_docx(X, F, path):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.95)
        sec.top_margin = sec.bottom_margin = Inches(0.85)
    grey = RGBColor(0x59, 0x59, 0x59)
    h = doc.add_paragraph()
    hr = h.add_run("Statistics Office  |  Section for Dwelling Statistics")
    hr.font.size = Pt(9)
    hr.font.color.rgb = grey
    foot = doc.sections[0].footer.paragraphs[0]
    fr = foot.add_run("Certification, 2027 vacant dwelling activation grant, %s pilot  |  Page " % W.COUNTY)
    fr.font.size = Pt(8)
    fr.font.color.rgb = grey
    _page_field(foot)
    sh, mi = W.PEOPLE["section_head"], W.PEOPLE["ministry"]
    addr = doc.add_paragraph()
    addr.add_run("%s\nHousing Policy Department\nAttn. %s\n\nOur ref. 26/4470-3   Your ref. letter of 1 December 2026"
                 "\n12 January 2027" % (W.MINISTRY, mi[0])).font.size = Pt(9.5)
    t = doc.add_paragraph()
    rt = t.add_run("Certification under Section 7: vacant dwellings in the %s pilot of the 2027 vacant dwelling "
                   "activation grant" % W.COUNTY)
    rt.bold = True
    rt.font.size = Pt(13)
    rt.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    qn_ = ["%s (%s per cent)" % (r["name"], r["share"]) for r in F["qualifying"]]
    p = doc.add_paragraph()
    p.add_run("%s municipalities qualify for the grant: %s. " % (F["count_word"], _join(qn_))).bold = True
    p.add_run("In each of them %s per cent or more of the dwellings the regulation counts (Section 3) stood vacant in "
              "the reference fortnight of 9 to 22 November 2026 (Sections 4 and 6). The highest share below the "
              "threshold is %s's %s per cent. Across the county %s of the %s counted dwellings stood vacant." % (
                  F["threshold"], F["highest_below"]["name"], F["highest_below"]["share"], F["totals"]["vacant"],
                  F["totals"]["counted"]))
    _table(doc, ["Municipality", "No.", "Counted dwellings", "Vacant dwellings", "Vacancy share (%)", "Qualifies"],
           [[m["name"], m["nr"], m["counted"], m["vacant"], m["share"], "Yes" if m["qualifies"] else "No"]
            for m in F["municipalities"]] + [["County", "", F["totals"]["counted"], F["totals"]["vacant"], "", ""]])
    doc.add_paragraph()
    doc.add_picture(io.BytesIO(chart_png(F)), width=Inches(6.2))
    cap = doc.add_paragraph()
    rc = cap.add_run("Figure 1. Vacancy share by municipality against the Section 6 threshold. Qualifying "
                     "municipalities in blue.")
    rc.italic = True
    rc.font.size = Pt(8)
    doc.add_paragraph(
        "How vacancy was established. Each counted dwelling was joined to its household metering points in the hub's "
        "extract, by the cadastre's unit ID or, for a house without one, by its address, and its hourly values were "
        "put on Norwegian time, the measured value kept where an hour had been estimated first, a dwelling's two "
        "points added where its meter was changed in the fortnight. A dwelling was taken as lived in when, on at "
        "least one day, its use in the morning or the evening rose above its own night-time use by more than %s kWh "
        "an hour, the household's day; it was vacant when no day showed one. The level of use does not settle it in "
        "November: an empty dwelling here is left heated, and the vacant dwellings used a median %s kWh a day against "
        "%s for the lived-in ones. Nor does the register: %s of the vacant dwellings still had someone registered in "
        "them on 1 November, and %s lived-in dwellings had nobody registered." % (
            F["cut"], F["vacant_median_kwh"], F["lived_median_kwh"], F["vacant_registered"], F["lived_unregistered"]))
    doc.add_paragraph("The dwelling-level list is in %s." % W.GOLDEN["xlsx"])
    a = doc.add_paragraph()
    a.add_run("Table A. Persons registered as resident in the counted dwellings on 1 November 2026, by age that "
              "day").bold = True
    _table(doc, ["Municipality", "No.", "Under 18", "18 to 66", "67 and over"],
           [[r["name"], r["nr"], r["under_18"], r["age_18_66"], r["age_67"]] for r in F["ask_a"]])
    n1 = doc.add_paragraph()
    r1 = n1.add_run("The register of 5 January 2027 with every event after 1 November undone. Where the cadastre "
                    "renumbered a block's flats in November the register shows each resident moving to the new unit "
                    "number on the day; those entries are a change of number, not a move.")
    r1.font.size = Pt(8.5)
    r1.font.color.rgb = grey
    b = doc.add_paragraph()
    b.add_run("Table B. Counted dwellings held on the land register on 1 November 2026 by the municipality or a company "
              "it owns").bold = True
    _table(doc, ["Municipality", "No.", "Dwellings"], [[r["name"], r["nr"], r["held"]] for r in F["ask_b"]])
    n2 = doc.add_paragraph()
    r2 = n2.add_run("Titles as registered on 1 November 2026, transfers registered after that day undone. A company the "
                    "municipality owns as defined in the ministry's letter, from the business register and the "
                    "shareholder register; titles from before the business register in the municipality's own name "
                    "are the municipality's.")
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = grey
    sig = doc.add_paragraph()
    sig.add_run("\nYours sincerely\n\n%s\n%s" % (sh[0], sh[1])).font.size = Pt(10.5)
    doc.save(path)


def main(target, golden, figures_path):
    import postprocess as PP
    os.makedirs(golden, exist_ok=True)
    X = compute(target)
    F = figures(X)
    xp = os.path.join(golden, W.GOLDEN["xlsx"])
    write_xlsx(X, F, xp)
    PP.normalize_ooxml(xp, W.PEOPLE["section_head"][0], W.PEOPLE["section_head"][0], "2027-01-08T09:14:00Z",
                       "2027-01-12T13:52:31Z", excel=True, company="Statistics Office")
    dp = os.path.join(golden, W.GOLDEN["docx"])
    write_docx(X, F, dp)
    PP.normalize_ooxml(dp, W.PEOPLE["section_head"][0], W.PEOPLE["section_head"][0], "2027-01-11T08:03:00Z",
                       "2027-01-12T14:20:00Z", company="Statistics Office")
    stamp = 1799761200.0
    for f in os.listdir(golden):
        os.utime(os.path.join(golden, f), (stamp, stamp))
    with open(figures_path, "w") as fh:
        json.dump(F, fh, indent=1, ensure_ascii=False)
        fh.write("\n")
    return X
