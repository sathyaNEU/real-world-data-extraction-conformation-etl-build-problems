"""Independent verifier: reads only the shipped bundle (its own readers: the csv module, pdfplumber, openpyxl, its
own time zone rule and tower-hours parser), takes the lead factors, the unrecorded operation's lead and the round's
threshold from the documents, rebuilds the 2026 inventory and checks the golden CSV and memo against it."""
import csv
import datetime as dt
import os
import re
import sys

FILES = dict(events="adsb_field_events_20251230_20270102.csv", master="MASTER.txt", dereg="DEREG.txt",
             acftref="ACFTREF.txt", based="state_based_aircraft_validation_2027.xlsx",
             atct="atct_daily_operations_2026.csv", f37="F37_2026_operations_count_report.pdf",
             l08="L08_2026_operations_count.xlsx", g42="G42_2026_activity_count.pdf", d19="D19_2026_daily_counts.xlsx",
             method="cvaqd_aircraft_lead_inventory_method_rev3.pdf",
             grants="cvaqd_2027_unleaded_avgas_transition_grants.pdf", network="cvaqd_adsb_receiver_network_notes.pdf",
             directory="state_airport_directory_2026.xlsx")
GOLD_CSV = "corrisan_2026_field_piston_lead.csv"
GOLD_MEMO = "cvaqd_2027_unleaded_grant_qualification.docx"


def pdf(p):
    import pdfplumber
    with pdfplumber.open(p) as d:
        return "\n".join(pg.extract_text() or "" for pg in d.pages)


def local(utc):
    # Pacific time: daylight time from 8 March 2026 10:00 UTC to 1 November 2026 09:00 UTC (the network notes)
    off = 7 if dt.datetime(2026, 3, 8, 10) <= utc < dt.datetime(2026, 11, 1, 9) else 8
    return utc - dt.timedelta(hours=off)


def fixed(p):
    with open(p, encoding="utf-8") as fh:
        rd = csv.reader(fh)
        head = [h.strip() for h in next(rd)]
        return [{h: v.strip() for h, v in zip(head, r)} for r in rd if r]


def d8(s):
    return dt.date(int(s[0:4]), int(s[4:6]), int(s[6:8]))


def main(target, golden):
    import openpyxl
    t = lambda k: os.path.join(target, FILES[k])
    ok = True

    def say(label, cond, extra=""):
        nonlocal ok
        ok &= bool(cond)
        print("verify %s  %s%s" % ("ok  " if cond else "FAIL", label, ("  [%s]" % extra) if not cond else ""))

    meth = pdf(t("method"))
    factors = {}
    for eng, row in (("single", r"One \(single-engine\)\s+([\d.]+) g\s+([\d.]+) g\s+([\d.]+) g"),
                     ("multi", r"Two or more \(multi-engine\)\s+([\d.]+) g\s+([\d.]+) g\s+([\d.]+) g")):
        m = re.search(row, meth)
        factors[(eng, "takeoff")], factors[(eng, "landing")], factors[(eng, "touch_and_go")] = map(float, m.groups())
    unrec_g = float(re.search(r"did not record is assigned ([\d.]+) g", meth).group(1))
    gr = pdf(t("grants"))
    threshold = float(re.search(r"is ([\d.]+) kg or more", gr).group(1))
    say("documents: Table 1 %s, %.3f g an unrecorded operation, threshold %.1f kg" % (factors, unrec_g, threshold),
        len(factors) == 6 and unrec_g > 0 and threshold > 0)

    # the register: every registration with its dates
    ref = {r["CODE"]: r for r in fixed(t("acftref"))}
    regs = []
    for r in fixed(t("master")):
        regs.append(("N" + r["N-NUMBER"], r["MODE S CODE HEX"].lower(), r["MFR MDL CODE"], r["SERIAL NUMBER"],
                     d8(r["CERT ISSUE DATE"]), None))
    for r in fixed(t("dereg")):
        regs.append(("N" + r["N-NUMBER"], r["MODE S CODE HEX"].lower(), r["MFR-MDL-CODE"], r["SERIAL-NUMBER"],
                     d8(r["CERT-ISSUE-DATE"]), d8(r["CANCEL-DATE"])))

    def engines(code):
        a = ref[code]
        if a["TYPE-ENG"] != "1" or a["TYPE-ACFT"] not in ("4", "5"):
            return None
        return "single" if a["NO-ENG"] == "1" else "multi"

    by_hex = {}
    for g in regs:
        by_hex.setdefault(g[1], []).append(g)

    def aircraft(h, day):
        gs = by_hex.get(h, [])
        for g in gs:
            if g[5] is not None and g[4] <= day < g[5]:
                return g
        cur = [g for g in gs if g[5] is None]
        return cur[0] if cur else None

    # the tower hours
    wb = openpyxl.load_workbook(t("directory"), read_only=True, data_only=True)
    rows = list(wb["Airports"].iter_rows(values_only=True))
    dirx = {r[0]: dict(zip(rows[0], r)) for r in rows[1:]}
    towers = sorted(l for l, d in dirx.items() if d["ATCT"] == "Y")
    cands = sorted(l for l, d in dirx.items() if d["ATCT"] == "N" and d["Use"] == "Public")
    hours = {}
    for l in towers:
        txt = dirx[l]["ATCT Hours (local)"]
        a, b = re.match(r"(\d\d)(\d\d)-(\d\d)(\d\d)", txt).group(1, 3), None
        main_ = [int(x) for x in re.match(r"(\d\d)(\d\d)-(\d\d)(\d\d)", txt).groups()]
        w = re.search(r"Nov-Feb (\d\d)(\d\d)-(\d\d)(\d\d)", txt)
        hours[l] = (main_, [int(x) for x in w.groups()] if w else None, "closed Dec 25 and Jan 1" in txt)

    def open_(l, ts):
        main_, w, closed = hours[l]
        if closed and (ts.month, ts.day) in ((12, 25), (1, 1)):
            return False
        hh = w if (w and ts.month in (11, 12, 1, 2)) else main_
        m = ts.hour * 60 + ts.minute
        return hh[0] * 60 + hh[1] <= m < hh[2] * 60 + hh[3]

    # the network's events
    lead = {l: 0.0 for l in dirx}
    pops = {l: 0 for l in dirx}
    heard_hours = {l: 0 for l in towers}
    heard = set()
    with open(t("events"), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            ts = local(dt.datetime.strptime(r["event_time_utc"], "%Y-%m-%dT%H:%M:%SZ"))
            if ts.year != 2026:
                continue
            l, h, e = r["site"], r["icao24"], r["event"]
            heard.add(h)
            n = 2 if e == "touch_and_go" else 1
            if l in hours and open_(l, ts):
                heard_hours[l] += n
            g = aircraft(h, ts.date())
            c = engines(g[2]) if g else None
            if c:
                lead[l] += factors[(c, e)] * n
                pops[l] += n
    # the counted fields
    unrec = {}
    with open(t("atct"), encoding="utf-8") as fh:
        civ = {}
        for r in csv.DictReader(fh):
            civ[r["Facility"]] = civ.get(r["Facility"], 0) + sum(int(r[c]) for c in (
                "Itinerant Air Carrier", "Itinerant Air Taxi", "Itinerant General Aviation", "Local Civil"))
    for l in towers:
        unrec[l] = civ[l] - heard_hours[l]
    study = {}
    m = re.search(r"\n2026 ([\d,]+) ", pdf(t("f37")))
    study["F37"] = int(m.group(1).replace(",", ""))
    m = re.search(r"\nYear ([\d,]+) ([\d,]+) ", pdf(t("g42")))
    study["G42"] = int(m.group(2).replace(",", ""))
    ws = openpyxl.load_workbook(t("l08"), read_only=True)["2026 by month"]
    study["L08"] = sum(r[1] + r[2] + 2 * r[3] for r in ws.iter_rows(min_row=5, max_row=16, values_only=True))
    ws = openpyxl.load_workbook(t("d19"), read_only=True)["Daily"]
    study["D19"] = sum(r[1] for r in ws.iter_rows(min_row=2, values_only=True) if r[0])
    for l, v in study.items():
        unrec[l] = v - pops[l]
    say("counted fields: %d towers and %d count studies" % (len(towers), len(study)), len(unrec) == 10)

    # the based piston airplanes the network never recorded (every address the airframe carried in 2026)
    by_n, by_air = {}, {}
    for g in regs:
        by_n.setdefault(g[0], []).append(g)
        by_air.setdefault((g[3], g[2]), []).append(g)
    nd = {l: 0 for l in dirx}
    wb = openpyxl.load_workbook(t("based"), read_only=True, data_only=True)
    for sh in wb.sheetnames:
        if sh == "About":
            continue
        for r in wb[sh].iter_rows(min_row=2, values_only=True):
            if not r[0] or r[0] not in nd:
                continue
            n = re.sub(r"[\s-]", "", str(r[2])).upper()
            n = n if n.startswith("N") else "N" + n
            gs = by_n[n]
            g = [x for x in gs if x[5] is None] or sorted(gs, key=lambda x: x[5])[-1:]
            g = g[0]
            if not engines(g[2]):
                continue
            hs = {g[1]} | {x[1] for x in by_air[(g[3], g[2])] if x[5] and x[5].year >= 2026}
            if not hs & heard:
                nd[r[0]] += 1
    k = sum(unrec.values()) / sum(nd[l] for l in unrec)
    say("k: %d unrecorded operations over %d never-recorded based piston airplanes, %.2f an airplane" % (
        sum(unrec.values()), sum(nd[l] for l in unrec), k), 60 < k < 80)
    inv = {}
    for l in dirx:
        if dirx[l]["Use"] != "Public":
            continue
        u = unrec[l] if l in unrec else k * nd[l]
        inv[l] = ((lead[l] + u * unrec_g) / 1000.0, pops[l] + u)
    q = sorted(l for l in cands if inv[l][0] >= threshold)
    with open(os.path.join(golden, GOLD_CSV), encoding="utf-8") as fh:
        gcsv = {r["loc_id"]: int(r["lead_2026_kg"]) for r in csv.DictReader(fh)}
    say("golden CSV: the 24 fields without a tower, each lead equal to the verifier's to the whole kilogram",
        sorted(gcsv) == cands and all(gcsv[l] == round(inv[l][0]) for l in cands),
        [l for l in cands if gcsv.get(l) != round(inv[l][0])])
    from docx import Document
    d = Document(os.path.join(golden, GOLD_MEMO))
    body = "\n".join(p.text for p in d.paragraphs)
    named = sorted(set(re.findall(r"\(([A-Z0-9]{3}), \d+ kg\)", body)))
    say("golden memo: the qualifying fields it names are the verifier's %s" % " ".join(q), named == q, named)
    tab = {r.cells[1].text: (r.cells[2].text, r.cells[3].text) for r in d.tables[0].rows[1:]}
    exp = {l: ("{:,}".format(int(round(inv[l][1] / 100.0)) * 100), str(round(inv[l][0]))) for l in towers}
    say("golden memo: the towered fields' operations and lead equal the verifier's", tab == exp, (tab, exp))
    print("verify: %s" % ("all checks pass" if ok else "FAILED"))
    if not ok:
        raise SystemExit("verify failed")
    return ok


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
