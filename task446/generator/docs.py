"""Documents and workbooks for task446's pack (architecture 3)."""
import datetime as dt
import io
import os
import re
import zipfile

import numpy as np
import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

import world
from world import D, EXTRACT_DATE

COMPANY = "Halstow Bank plc"
ORDER = ["[Content_Types].xml", "_rels/.rels", "xl/workbook.xml", "xl/_rels/workbook.xml.rels", "word/document.xml",
         "word/_rels/document.xml.rels"]


def _iso(t):
    return t.strftime("%Y-%m-%dT%H:%M:%SZ")


def core_xml(creator, last_by, created, modified, revision=None):
    rev = "<cp:revision>%d</cp:revision>" % revision if revision else ""
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
            '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
            'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
            'xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
            '<dc:title></dc:title><dc:subject></dc:subject><dc:creator>%s</dc:creator><cp:keywords></cp:keywords>'
            '<dc:description></dc:description><cp:lastModifiedBy>%s</cp:lastModifiedBy>%s'
            '<dcterms:created xsi:type="dcterms:W3CDTF">%s</dcterms:created>'
            '<dcterms:modified xsi:type="dcterms:W3CDTF">%s</dcterms:modified></cp:coreProperties>'
            % (creator, last_by, rev, _iso(created), _iso(modified))).encode()


def excel_app_xml(sheets):
    parts = "".join("<vt:lpstr>%s</vt:lpstr>" % html_escape(t) for t in sheets)
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
            '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
            'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
            '<Application>Microsoft Excel</Application><DocSecurity>0</DocSecurity><ScaleCrop>false</ScaleCrop>'
            '<HeadingPairs><vt:vector size="2" baseType="variant"><vt:variant><vt:lpstr>Worksheets</vt:lpstr>'
            '</vt:variant><vt:variant><vt:i4>%d</vt:i4></vt:variant></vt:vector></HeadingPairs>'
            '<TitlesOfParts><vt:vector size="%d" baseType="lpstr">%s</vt:vector></TitlesOfParts>'
            '<Company>%s</Company><LinksUpToDate>false</LinksUpToDate><SharedDoc>false</SharedDoc>'
            '<HyperlinksChanged>false</HyperlinksChanged><AppVersion>16.0300</AppVersion></Properties>'
            % (len(sheets), len(sheets), parts, COMPANY)).encode()


def html_escape(t):
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def write_container(path, members):
    """Write an OOXML package the way Office does: fixed member order, 1980 member times."""
    tail = ["docProps/core.xml", "docProps/app.xml"]
    names = [n for n in ORDER if n in members] + sorted(n for n in members if n not in ORDER + tail) + \
        [n for n in tail if n in members]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n in names:
            zi = zipfile.ZipInfo(n, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, members[n])


def finish_xlsx(path, creator, last_by, created, modified):
    with zipfile.ZipFile(path) as z:
        members = {i.filename: z.read(i.filename) for i in z.infolist()}
    wb = openpyxl.load_workbook(path, read_only=True)
    sheets = wb.sheetnames
    wb.close()
    members["docProps/app.xml"] = excel_app_xml(sheets)
    members["docProps/core.xml"] = core_xml(creator, last_by, created, modified)
    write_container(path, members)


def widths(ws, ws_widths):
    for i, w_ in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w_


# ------------------------------------------------------------------ group register
def register_workbook(w, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Group register"
    ws["A1"] = "Credit system: obligor group register"
    ws["A1"].font = Font(bold=True)
    ws["A2"] = "As at 30/09/2026. Reference Data quarter-end extract. Includes groups whose relationship has ended."
    hdr = ["group_id", "group_name", "ultimate_parent_lei", "relationship_start", "relationship_end", "sector_code",
           "sector", "coverage_team", "credit_officer"]
    ws.append([])
    ws.append(hdr)
    for c in ws[4]:
        c.font = Font(bold=True)
    for gid, g in sorted(w.groups.items()):
        top = g.members[0]
        code, desc, _ = world.SECTORS[g.sector]
        ws.append([gid, g.name, top.lei or "", g.active_from, g.active_to, code, desc, g.coverage, g.officer])
    for row in ws.iter_rows(min_row=5, min_col=4, max_col=5):
        for c in row:
            c.number_format = "dd/mm/yyyy"
    widths(ws, [11, 44, 23, 15, 15, 11, 36, 24, 13])
    ws.freeze_panes = "A5"
    wb.save(path)
    t = dt.datetime(2026, 10, 1, 5, 5, 0)
    finish_xlsx(path, "Reference Data", "Reference Data", t, t)


def ratings_workbook(w, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master scale"
    ws.append(["Halstow Bank internal rating master scale (corporate and commercial), effective 1 January 2026"])
    ws["A1"].font = Font(bold=True)
    ws.append([])
    ws.append(["grade", "description", "pd_lower", "pd_upper", "pd_midpoint", "s&p_equivalent"])
    for c in ws[3]:
        c.font = Font(bold=True)
    scale = [(1, "Exceptional", 0.0000, 0.0005, "AA"), (2, "Very strong", 0.0005, 0.0012, "A"),
             (3, "Strong", 0.0012, 0.0025, "BBB+"), (4, "Good", 0.0025, 0.0050, "BBB-"),
             (5, "Satisfactory", 0.0050, 0.0100, "BB+"), (6, "Adequate", 0.0100, 0.0200, "BB"),
             (7, "Weak", 0.0200, 0.0450, "BB-"), (8, "Vulnerable", 0.0450, 0.1000, "B"),
             (9, "Watch", 0.1000, 0.2500, "CCC"), (10, "Default", 1.0, 1.0, "D")]
    for gr, de, lo, hi, sp in scale:
        mid = round((lo * hi) ** 0.5, 5) if lo > 0 and hi < 1 else (0.0002 if gr == 1 else 1.0)
        ws.append([gr, de, lo, hi, mid, sp])
    widths(ws, [8, 16, 10, 10, 12, 16])
    ws2 = wb.create_sheet("Group ratings 30-09-2026")
    ws2.append(["group_id", "rating_grade", "rating_date", "rating_model"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    rng = np.random.default_rng(world.SEED + 31)
    for gid, g in sorted(w.groups.items()):
        if not w.active(g, EXTRACT_DATE):
            continue
        base = 4.2 + (1.8 if g.flag.get(4) else 0.0) + (2.5 if w.on_list(g, EXTRACT_DATE) else 0.0)
        grade = int(np.clip(round(base + rng.normal(0, 1.1)), 1, 9))
        rdate = EXTRACT_DATE - dt.timedelta(days=int(rng.integers(10, 360)))
        ws2.append([gid, grade, rdate, "CORP-SCORE v5" if rng.random() < 0.8 else "COMMERCIAL-SCORE v3"])
    for row in ws2.iter_rows(min_row=2, min_col=3, max_col=3):
        for c in row:
            c.number_format = "dd/mm/yyyy"
    widths(ws2, [11, 13, 13, 22])
    wb.save(path)
    finish_xlsx(path, "Credit Risk Models", "Credit Risk Models", dt.datetime(2025, 12, 12, 10, 21, 0),
                dt.datetime(2026, 10, 2, 8, 47, 0))


def outlook_workbook(path):
    wb = openpyxl.Workbook()
    qs = ["2026 Q4", "2027 Q1", "2027 Q2", "2027 Q3", "2027 Q4", "2028 Q1", "2028 Q2", "2028 Q3", "2028 Q4", "2029 Q1",
          "2029 Q2", "2029 Q3", "2029 Q4"]
    paths = {
        "Central": dict(gdp=[0.2, 0.3, 0.3, 0.4, 0.4, 0.4, 0.4, 0.4, 0.4, 0.4, 0.4, 0.4, 0.4],
                        unemp=[5.1, 5.2, 5.2, 5.1, 5.0, 4.9, 4.9, 4.8, 4.8, 4.7, 4.7, 4.7, 4.6],
                        rate=[3.75, 3.75, 3.50, 3.50, 3.25, 3.25, 3.25, 3.00, 3.00, 3.00, 3.00, 3.00, 3.00],
                        cpi=[3.4, 3.1, 2.8, 2.6, 2.4, 2.2, 2.1, 2.0, 2.0, 2.0, 2.0, 2.0, 2.0]),
        "Downside": dict(gdp=[-0.3, -0.4, -0.2, 0.0, 0.2, 0.3, 0.3, 0.4, 0.4, 0.4, 0.4, 0.4, 0.4],
                         unemp=[5.4, 5.8, 6.1, 6.3, 6.4, 6.3, 6.2, 6.0, 5.9, 5.8, 5.7, 5.6, 5.5],
                         rate=[3.75, 3.50, 3.00, 2.75, 2.50, 2.50, 2.50, 2.50, 2.75, 2.75, 2.75, 3.00, 3.00],
                         cpi=[3.6, 3.5, 3.1, 2.7, 2.3, 2.0, 1.8, 1.7, 1.8, 1.9, 2.0, 2.0, 2.0]),
        "Severe downside": dict(gdp=[-0.9, -1.2, -0.8, -0.3, 0.1, 0.3, 0.4, 0.4, 0.5, 0.5, 0.5, 0.5, 0.5],
                                unemp=[5.8, 6.6, 7.3, 7.8, 8.0, 7.9, 7.7, 7.4, 7.1, 6.9, 6.7, 6.5, 6.3],
                                rate=[3.75, 3.25, 2.50, 2.00, 1.75, 1.75, 1.75, 2.00, 2.00, 2.25, 2.25, 2.50, 2.50],
                                cpi=[3.8, 3.9, 3.4, 2.6, 1.9, 1.4, 1.2, 1.3, 1.5, 1.7, 1.9, 2.0, 2.0]),
    }
    first = True
    for name, p in paths.items():
        s = wb.active if first else wb.create_sheet(name)
        s.title = name
        first = False
        s["A1"] = "Group Economics: UK outlook, October 2026. %s path." % name
        s["A1"].font = Font(bold=True)
        s.append([])
        s.append(["Quarter", "Real GDP, % q/q", "Unemployment rate, %", "Bank Rate, % (end quarter)",
                  "CPI inflation, % y/y"])
        for c in s[3]:
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, vertical="bottom")
        for i, q in enumerate(qs):
            s.append([q, p["gdp"][i], p["unemp"][i], p["rate"][i], p["cpi"][i]])
            for col in (2, 3, 5):
                s.cell(row=s.max_row, column=col).number_format = "0.0"
            s.cell(row=s.max_row, column=4).number_format = "0.00"
        widths(s, [11, 14, 15, 17, 15])
        s.freeze_panes = "B4"
    n = wb.create_sheet("Notes")
    for row in [
        ["Notes"],
        ["Prepared by Group Economics for planning and stress-testing users across the Group."],
        ["Paths are quarterly. GDP growth is quarter on quarter; unemployment is the quarterly average (ILO, 16+);"],
        ["Bank Rate is the rate at quarter end; CPI is the quarterly average, year on year."],
        ["Central reflects the Group's view at 1 October 2026. Downside and Severe downside are narrative paths for"],
        ["sensitivity analysis and carry no probability."],
        ["Next update: January 2027."],
    ]:
        n.append(row)
    n["A1"].font = Font(bold=True)
    widths(n, [110])
    wb.save(path)
    finish_xlsx(path, "Group Economics", "Group Economics", dt.datetime(2026, 10, 5, 13, 12, 0),
                dt.datetime(2026, 10, 6, 9, 38, 0))


# ------------------------------------------------------------------ docx agenda
W_NS = ('xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:o="urn:schemas-microsoft-com:office:office" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:v="urn:schemas-microsoft-com:vml" '
        'xmlns:w10="urn:schemas-microsoft-com:office:word" '
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" '
        'xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml" '
        'xmlns:wne="http://schemas.microsoft.com/office/word/2006/wordml" mc:Ignorable="w14 w15"')
XML_HEAD = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'


def _w_run(text, bold=False, italic=False):
    props = ("<w:b/>" if bold else "") + ("<w:i/>" if italic else "")
    rpr = "<w:rPr>%s</w:rPr>" % props if props else ""
    parts = []
    for k, piece in enumerate(text.split("\t")):
        if k:
            parts.append("<w:tab/>")
        if piece:
            sp = ' xml:space="preserve"' if piece != piece.strip() else ""
            parts.append("<w:t%s>%s</w:t>" % (sp, html_escape(piece)))
    return "<w:r>%s%s</w:r>" % (rpr, "".join(parts))


def _w_par(runs, rsid, style=None, ind=None, after=None):
    ppr = ""
    if style or ind or after is not None:
        ppr = "<w:pPr>%s%s%s</w:pPr>" % ('<w:pStyle w:val="%s"/>' % style if style else "",
                                         '<w:spacing w:after="%d"/>' % after if after is not None else "",
                                         '<w:ind w:left="%d"/>' % ind if ind else "")
    return '<w:p w:rsidR="%s" w:rsidRDefault="%s">%s%s</w:p>' % (rsid, rsid, ppr, "".join(runs))


def agenda_docx(path):
    rs = ["00A31C58", "00A31C58", "00A31C58", "00A31C58", "00F27D14", "00A31C58", "00F27D14", "00A31C58", "00A31C58",
          "002B6E90", "00F27D14", "00A31C58", "00A31C58", "002B6E90", "00A31C58"]
    paras = [
        ([_w_run("HALSTOW BANK PLC", bold=True)], None, None, 0),
        ([_w_run("Executive Risk Committee", bold=True)], None, None, 0),
        ([_w_run("Wednesday 4 November 2026, 09:30 to 12:00, Board Room 2, Halstow House")], None, None, None),
        ([_w_run("Chair: Chief Risk Officer. Members: Chief Financial Officer, Chief Credit Officer, Chief Operating "
                 "Officer. In attendance: Jill Berry (Head of Credit Risk Monitoring), Mathew Long (Head of "
                 "Restructuring), Head of Business Support, Company Secretary.")], None, None, None),
        ([_w_run("Agenda", bold=True)], None, None, None),
        ([_w_run("1.\tMinutes of the meeting of 7 October 2026 and matters arising")], None, None, None),
        ([_w_run("2.\tQ3 2026 risk appetite dashboard (Chief Risk Officer)")], None, None, None),
        ([_w_run("3.\t2027 resourcing budget: credit")], None, None, 60),
        ([_w_run("3.1\tWatch-list caseload for the year to 30 September 2027 (Credit Risk Monitoring, paper to "
                 "follow). The caseload standard for Restructuring and Business Support is twelve watch-list groups "
                 "per analyst.")], None, 720, 60),
        ([_w_run("3.2\tRestructuring and Business Support headcount for 2027 (Head of Restructuring)")], None, 720, None),
        ([_w_run("4.\tEarly-warning system annual performance report: taken at the Credit Risk Committee on 18 "
                 "November")], None, None, None),
        ([_w_run("5.\tOperational resilience: third-party data feeds (Chief Operating Officer)")], None, None, None),
        ([_w_run("6.\tAny other business")], None, None, None),
        ([_w_run("Papers to the Company Secretary by 12:00 on Wednesday 28 October. Item 3.1 papers carry the "
                 "committed exposure by caseload band so that 3.2 can be read against them.")], None, None, None),
        ([_w_run("Internal. Not for onward distribution.", italic=True)], None, None, None),
    ]
    body = "".join(_w_par(r, rs[k], style, ind, after) for k, (r, style, ind, after) in enumerate(paras))
    sect = ('<w:sectPr w:rsidR="00A31C58"><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1440" w:right="1440" '
            'w:bottom="1440" w:left="1440" w:header="708" w:footer="708" w:gutter="0"/><w:cols w:space="708"/>'
            '<w:docGrid w:linePitch="360"/></w:sectPr>')
    document = XML_HEAD + '<w:document %s><w:body>%s%s</w:body></w:document>' % (W_NS, body, sect)
    styles = (XML_HEAD + '<w:styles %s><w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:ascii="Arial" '
              'w:eastAsiaTheme="minorHAnsi" w:hAnsi="Arial" w:cs="Arial"/><w:sz w:val="20"/><w:szCs w:val="20"/>'
              '<w:lang w:val="en-GB" w:eastAsia="en-US" w:bidi="ar-SA"/></w:rPr></w:rPrDefault><w:pPrDefault><w:pPr>'
              '<w:spacing w:after="160" w:line="259" w:lineRule="auto"/></w:pPr></w:pPrDefault></w:docDefaults>'
              '<w:latentStyles w:defLockedState="0" w:defUIPriority="99" w:defSemiHidden="0" w:defUnhideWhenUsed="0" '
              'w:defQFormat="0" w:count="376"><w:lsdException w:name="Normal" w:uiPriority="0" w:qFormat="1"/>'
              '<w:lsdException w:name="heading 1" w:uiPriority="9" w:qFormat="1"/></w:latentStyles>'
              '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/><w:qFormat/>'
              '</w:style><w:style w:type="character" w:default="1" w:styleId="DefaultParagraphFont"><w:name '
              'w:val="Default Paragraph Font"/><w:uiPriority w:val="1"/><w:semiHidden/><w:unhideWhenUsed/></w:style>'
              '<w:style w:type="table" w:default="1" w:styleId="TableNormal"><w:name w:val="Normal Table"/>'
              '<w:uiPriority w:val="99"/><w:semiHidden/><w:unhideWhenUsed/><w:tblPr><w:tblInd w:w="0" w:type="dxa"/>'
              '<w:tblCellMar><w:top w:w="0" w:type="dxa"/><w:left w:w="108" w:type="dxa"/><w:bottom w:w="0" '
              'w:type="dxa"/><w:right w:w="108" w:type="dxa"/></w:tblCellMar></w:tblPr></w:style><w:style '
              'w:type="numbering" w:default="1" w:styleId="NoList"><w:name w:val="No List"/><w:uiPriority w:val="99"/>'
              '<w:semiHidden/><w:unhideWhenUsed/></w:style></w:styles>' % W_NS)
    settings = (XML_HEAD + '<w:settings %s><w:zoom w:percent="110"/><w:proofState w:spelling="clean" '
                'w:grammar="clean"/><w:defaultTabStop w:val="720"/><w:characterSpacingControl w:val="doNotCompress"/>'
                '<w:compat><w:compatSetting w:name="compatibilityMode" w:uri="http://schemas.microsoft.com/office/word"'
                ' w:val="15"/><w:compatSetting w:name="overrideTableStyleFontSizeAndJustification" '
                'w:uri="http://schemas.microsoft.com/office/word" w:val="1"/><w:compatSetting '
                'w:name="enableOpenTypeFeatures" w:uri="http://schemas.microsoft.com/office/word" w:val="1"/>'
                '<w:compatSetting w:name="doNotFlipMirrorIndents" w:uri="http://schemas.microsoft.com/office/word" '
                'w:val="1"/></w:compat><w:rsids><w:rsidRoot w:val="00A31C58"/><w:rsid w:val="002B6E90"/><w:rsid '
                'w:val="00A31C58"/><w:rsid w:val="00F27D14"/></w:rsids><m:mathPr><m:mathFont m:val="Cambria Math"/>'
                '</m:mathPr><w:themeFontLang w:val="en-GB"/><w:clrSchemeMapping w:bg1="light1" w:t1="dark1" '
                'w:bg2="light2" w:t2="dark2" w:accent1="accent1" w:accent2="accent2" w:accent3="accent3" '
                'w:accent4="accent4" w:accent5="accent5" w:accent6="accent6" w:hyperlink="hyperlink" '
                'w:followedHyperlink="followedHyperlink"/><w:shapeDefaults><o:shapedefaults v:ext="edit" '
                'spidmax="1026"/><o:shapelayout v:ext="edit"><o:idmap v:ext="edit" data="1"/></o:shapelayout>'
                '</w:shapeDefaults><w:decimalSymbol w:val="."/><w:listSeparator w:val=","/>'
                '<w14:docId w14:val="5E2A7C31"/><w15:chartTrackingRefBased/><w15:docId '
                'w15:val="{8C4B2F0A-6D1E-4F3B-9A57-2E61C0D4B7A9}"/></w:settings>' % W_NS)
    fonts = (XML_HEAD + '<w:fonts %s>' % W_NS + "".join(
        '<w:font w:name="%s"><w:panose1 w:val="%s"/><w:charset w:val="00"/><w:family w:val="%s"/>'
        '<w:pitch w:val="variable"/></w:font>' % x for x in
        [("Arial", "020B0604020202020204", "swiss"), ("Times New Roman", "02020603050405020304", "roman"),
         ("Calibri Light", "020F0302020204030204", "swiss"), ("Calibri", "020F0502020204030204", "swiss")])
             + '</w:fonts>')
    web = (XML_HEAD + '<w:webSettings %s><w:optimizeForBrowser/><w:allowPNG/></w:webSettings>' % W_NS)
    from openpyxl.writer.theme import theme_xml
    theme = XML_HEAD + re.sub(r">\s+<", "><", theme_xml.split("?>", 1)[1].strip())
    ct = (XML_HEAD + '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
          '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
          '<Default Extension="xml" ContentType="application/xml"/>'
          '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.'
          'wordprocessingml.document.main+xml"/><Override PartName="/word/styles.xml" ContentType="application/'
          'vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/><Override PartName="/word/settings.xml" '
          'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml"/><Override '
          'PartName="/word/webSettings.xml" ContentType="application/vnd.openxmlformats-officedocument.'
          'wordprocessingml.webSettings+xml"/><Override PartName="/word/fontTable.xml" ContentType="application/'
          'vnd.openxmlformats-officedocument.wordprocessingml.fontTable+xml"/><Override PartName="/word/theme/'
          'theme1.xml" ContentType="application/vnd.openxmlformats-officedocument.theme+xml"/><Override PartName='
          '"/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
          '<Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.'
          'extended-properties+xml"/></Types>')
    rels = (XML_HEAD + '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/'
            'core-properties" Target="docProps/core.xml"/><Relationship Id="rId2" Type="http://schemas.openxmlformats'
            '.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/><Relationship '
            'Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
            'Target="word/document.xml"/></Relationships>')
    drels = (XML_HEAD + '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
             '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
             'webSettings" Target="webSettings.xml"/><Relationship Id="rId2" Type="http://schemas.openxmlformats.org/'
             'officeDocument/2006/relationships/settings" Target="settings.xml"/><Relationship Id="rId1" Type="http://'
             'schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/><Relationship '
             'Id="rId5" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme" '
             'Target="theme/theme1.xml"/><Relationship Id="rId4" Type="http://schemas.openxmlformats.org/'
             'officeDocument/2006/relationships/fontTable" Target="fontTable.xml"/></Relationships>')
    text = " ".join(re.sub("<[^>]+>", " ", body).split())
    words = len(text.split())
    chars = len(text.replace(" ", ""))
    app = (XML_HEAD + '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
           'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"><Template>Normal.dotm'
           '</Template><TotalTime>38</TotalTime><Pages>1</Pages><Words>%d</Words><Characters>%d</Characters>'
           '<Application>Microsoft Office Word</Application><DocSecurity>0</DocSecurity><Lines>%d</Lines>'
           '<Paragraphs>%d</Paragraphs><ScaleCrop>false</ScaleCrop><Company>%s</Company><LinksUpToDate>false'
           '</LinksUpToDate><CharactersWithSpaces>%d</CharactersWithSpaces><SharedDoc>false</SharedDoc>'
           '<HyperlinksChanged>false</HyperlinksChanged><AppVersion>16.0000</AppVersion></Properties>'
           % (words, chars, 24, len(paras), COMPANY, len(text)))
    members = {
        "[Content_Types].xml": ct.encode(), "_rels/.rels": rels.encode(), "word/document.xml": document.encode(),
        "word/_rels/document.xml.rels": drels.encode(), "word/styles.xml": styles.encode(),
        "word/settings.xml": settings.encode(), "word/fontTable.xml": fonts.encode(),
        "word/webSettings.xml": web.encode(), "word/theme/theme1.xml": theme.encode(), "docProps/app.xml": app.encode(),
        "docProps/core.xml": core_xml("Company Secretariat", "Abbie Stevens", dt.datetime(2026, 10, 7, 9, 14, 0),
                                      dt.datetime(2026, 10, 8, 15, 52, 0), revision=4),
    }
    write_container(path, members)


# ------------------------------------------------------------------ PDF standard
def _pdf(path, title, author, blocks):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab import rl_config
    rl_config.invariant = 1
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=13)
    bullet = ParagraphStyle("u", parent=body, leftIndent=14)
    h = ParagraphStyle("h", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=10.5, leading=14,
                       spaceBefore=6)
    t = ParagraphStyle("t", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=13, leading=17)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=56, rightMargin=56, topMargin=56, bottomMargin=56,
                            title=title, author=author, creator="", producer="", subject="")
    story = []
    for kind, text in blocks:
        story.append(Paragraph(text, {"t": t, "h": h, "p": body, "u": bullet}[kind]))
        story.append(Spacer(1, 4 if kind != "t" else 8))
    doc.build(story)
    data = buf.getvalue()
    stamp = b"D:20260212101500+00'00'"
    data = data.replace(b"D:20000101000000+00'00'", stamp)
    for s in (b"ReportLab Generated PDF document (opensource)",
              b"ReportLab generated PDF document -- digest (opensource)"):
        data = data.replace(s, b" " * len(s))
    with open(path, "wb") as fh:
        fh.write(data)


def standard_pdf(path):
    B = [
        ("t", "Halstow Bank plc<br/>Early-Warning and Watch-List Monitoring Standard"),
        ("p", "Standard CRM-07, version 3.1. Approved by the Credit Risk Committee on 12 February 2026. "
              "Owner: Head of Credit Risk Monitoring. Next review: February 2027."),
        ("h", "1 Scope"),
        ("p", "1.1 This standard covers the corporate and commercial book: every obligor group with a committed "
              "facility from Halstow Bank plc or Halstow Bank Ireland DAC. It sets how early-warning signals are "
              "carried to groups, how the watch list is measured against them, and how the watch-list caseload "
              "is forecast for resourcing."),
        ("p", "1.2 The early-warning system is release 3.2 of the engine, in production since 10 January 2022. "
              "Its models, triggers and bands have not changed since. A change would be a new release, approved "
              "by the Model Risk Committee before use."),
        ("p", "1.3 Dates and times in this standard are UK time (GMT or BST, whichever is in force)."),
        ("h", "2 Obligor groups"),
        ("p", "2.1 An obligor group is the set of legal entities under one ultimate accounting-consolidation "
              "parent. The credit system's group register holds each group, the dates its relationship with the "
              "Bank started and ended, and the LEI of its ultimate parent where the parent has one."),
        ("p", "2.2 A group's members are its ultimate parent, the customers the credit system books to it and "
              "every legal entity whose GLEIF relationship record shows it accounting-consolidated into the "
              "ultimate parent "
              "(IS_ULTIMATELY_CONSOLIDATED_BY; where an entity reports only IS_DIRECTLY_CONSOLIDATED_BY, its "
              "direct parents are followed to the top), with the relationship in force at the date in question."),
        ("p", "2.3 Group structure is read from GLEIF's Level 1 and Level 2 data as GLEIF publishes it, under "
              "GLEIF's own definitions of its relationship types and registration statuses."),
        ("h", "3 Signals and flags"),
        ("p", "3.1 The engine publishes a signal for every entity it covers at each month-end: GREEN, AMBER or "
              "RED, with a score from 0 to 99 and the triggers behind it. Entities with an LEI are keyed by LEI, "
              "entities without one by their Companies House number."),
        ("p", "3.2 A group is flagged at a date when any of its members carries AMBER or RED in the engine's "
              "publication for that month-end."),
        ("h", "4 The watch list and the cohorts"),
        ("p", "4.1 The monitored book at a cohort date, 30 September each year, is every group active on the "
              "group register at that date and not on the watch list at that date."),
        ("p", "4.2 A group in a cohort reaches the watch list if it is placed on the watch list in the twelve "
              "months after the cohort date. A group that leaves the book in that time without being placed has "
              "not reached it. A group counts once. A cohort is closed once its twelve months have ended."),
        ("p", "4.3 The system's performance is reported to the Credit Risk Committee each November on the closed "
              "cohorts scored by the current release, taken together:"),
        ("u", "hit rate: groups flagged at the cohort date that reach the watch list, as a share of groups "
              "flagged;"),
        ("u", "capture rate: groups reaching the watch list that were flagged at the cohort date, as a share of "
              "groups reaching it;"),
        ("u", "false-alarm rate: groups not reaching the watch list that were flagged at the cohort date, as a "
              "share of groups not reaching it."),
        ("h", "5 Watch-list caseload forecast"),
        ("p", "5.1 Each October credit risk monitoring forecasts how many groups in the monitored book at the "
              "cohort date will reach the watch list in the following twelve months, and the committed exposure "
              "they will bring, for the resourcing budget. Groups that join the book after the cohort date are "
              "outside the forecast."),
        ("p", "5.2 The forecast is the number of groups in the monitored book at the cohort date expected to reach "
              "the watch list in the twelve months after it, with the book, the outcome and the groups left out as "
              "4.1, 4.2 and 5.1 set them, and the committed exposure at the cohort date that those groups are "
              "expected to bring. The Committee resources the coming year's caseload on it, and each forecast is "
              "set against its cohort's actual arrivals once the cohort closes."),
        ("p", "5.3 The caseload is reported to the nearest whole group, and the expected arrivals and the committed "
              "exposure they are expected to bring in each band of 6.2 to one decimal."),
        ("h", "6 Committed exposure and caseload bands"),
        ("p", "6.1 A group's committed exposure at a date is the sum of the committed limits, as they stood at the "
              "end of that date, of the facilities on which one or more of its members is an obligor, across "
              "Halstow Bank plc and Halstow Bank Ireland DAC, in sterling at that date's Group Treasury closing "
              "rates. Only facilities in force at the end of the date count: a facility whose start date falls "
              "after the date, or whose final maturity date falls on or before it, is not in force at the end of "
              "it."),
        ("p", "6.2 The caseload is reported in five bands of committed exposure at the cohort date: under GBP 5m; "
              "GBP 5m and under "
              "GBP 15m; GBP 15m and under GBP 40m; GBP 40m and under GBP 100m; GBP 100m and over. Groups of "
              "GBP 40m and over are worked by Restructuring, the rest by Business Support."),
        ("h", "7 Records"),
        ("p", "7.1 Signals are kept for every month-end from September 2022. Reference Data extracts the group "
              "register, the obligor master, the watch-list history and the GLEIF records for the Bank's entities "
              "at each quarter-end."),
    ]
    _pdf(path, "Early-Warning and Watch-List Monitoring Standard", "Credit Risk Monitoring", B)


# ------------------------------------------------------------------ markdown and text
def field_reference(path, files):
    txt = f"""# Credit data: field reference

Maintained by Reference Data. Covers the quarter-end extracts used by Credit Risk Monitoring.

## {files['signals']}

Month-end publications of the early-warning engine, one row per covered entity per month-end.

| field | meaning |
|---|---|
| publication_date | the month-end the publication is for |
| entity_key | the entity's LEI, or `CRN:` and the Companies House number for an entity without an LEI |
| entity_name | the name as the engine's data vendor holds it |
| country | country of registration |
| signal | GREEN, AMBER or RED |
| score | 0 to 99; GREEN below 40, AMBER 40 to 69, RED 70 and over |
| trigger_codes | triggers active at the publication, semicolon separated (see {files['triggers']}) |
| feed | the vendor feed the leading trigger came from; COMPOSITE where no trigger is active |
| engine_release | engine release that produced the publication |

## {files['obligors']}

Every customer on the credit system, open or closed.

| field | meaning |
|---|---|
| customer_id | credit system customer number |
| customer_name | the customer's registered name |
| lei | the customer's LEI where held |
| companies_house_no | Companies House number for UK-registered customers |
| country | country of registration |
| group_id | the obligor group the credit system books the customer to; blank until the monthly grouping review books a new customer |
| customer_since | the date the customer was opened on the credit system |
| customer_status | ACTIVE or CLOSED at the extract date |
| primary_booking_entity | HB-PLC (Halstow Bank plc) or HB-DAC (Halstow Bank Ireland DAC) |

## {files['register']}

One row per obligor group, including groups whose relationship has ended.

| field | meaning |
|---|---|
| ultimate_parent_lei | LEI of the group's ultimate accounting-consolidation parent; blank where the parent has none |
| relationship_start | date the group's first facility started |
| relationship_end | date the relationship ended (last facility repaid, refinanced elsewhere or transferred to Recoveries); blank while active |
| sector_code, sector | the group's sector |

## {files['watchlist']}

Every watch-list case opened since 1 January 2020.

| field | meaning |
|---|---|
| placed_on | date the group was placed on the watch list |
| watch_level | 1, 2 or 3 (case intensity); every level is on the watch list |
| removed_on | date the case closed; blank while open |
| removal_reason | CURED, REFINANCED_AWAY or TO_RECOVERY |

## {files['facilities']}

Facilities on the Halstow Bank plc and Halstow Bank Ireland DAC loan ledgers when the quarter-end extract ran, with
their obligors.

| field | meaning |
|---|---|
| booking_entity | HB-PLC or HB-DAC, the ledger the facility is booked on |
| facility_id | the facility number on that ledger |
| customer_id | the obligor the row is for |
| product | facility product |
| currency | facility currency |
| committed_limit | the facility's committed limit when the extract ran, in the facility currency |
| drawn_balance | the obligor's drawings when the extract ran, in the facility currency |
| start_date | date the facility became available |
| final_maturity_date | date the facility falls due for repayment and the commitment ends |
| facility_status | ACTIVE, CANCELLED or CLOSED when the extract ran |

## {files['limits']}

Every committed limit each facility has carried, one row per version.

| field | meaning |
|---|---|
| booking_entity, facility_id | as in {files['facilities']} |
| committed_limit | the committed limit of the version, in the facility currency |
| effective_from | when the version took effect, UTC, as the loan platform records it |
| source | SOD_BATCH for the start-of-day batch, otherwise the user who keyed it |

## {files['fx']}

| field | meaning |
|---|---|
| units_per_gbp | units of the currency for one pound sterling at the Group Treasury 16:00 London close |

## GLEIF files ({files['lei']}, {files['rr']}, {files['repex']})

Records from GLEIF's concatenated files published on 30 September 2026 (LEI-CDF 3.1, RR-CDF 2.1 and reporting
exceptions), filtered by Reference Data to the LEIs in the obligor master and the engine's coverage list.
Field names are GLEIF's.
"""
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(txt)


def count_rows(path):
    with open(path, encoding="utf-8") as fh:
        return sum(1 for _ in fh) - 1


def extract_log(path, files, target):
    def rows(k):
        return count_rows(os.path.join(target, files[k]))
    lines = [
        "REFERENCE DATA EXTRACT LOG",
        "Halstow Bank plc   quarter-end pack for Credit Risk Monitoring   as at 30/09/2026   compiled 08/10/2026 16:40",
        "",
        "job                          produced           rows     note",
        "EWS_MONTHEND_PUBLICATIONS    2026-10-02 06:14   %7d   publications 2022-09-30 to 2026-09-30" % rows("signals"),
        "GLEIF_LEI_CDF_SUBSET         2026-10-01 05:40   %7d   GLEIF concatenated file of 2026-09-30" % rows("lei"),
        "GLEIF_RR_CDF_SUBSET          2026-10-01 05:41   %7d   GLEIF concatenated file of 2026-09-30" % rows("rr"),
        "GLEIF_REPEX_SUBSET           2026-10-01 05:41   %7d   GLEIF concatenated file of 2026-09-30" % rows("repex"),
        "CREDIT_OBLIGOR_MASTER        2026-10-01 06:02   %7d" % rows("obligors"),
        "CREDIT_GROUP_REGISTER        2026-10-01 06:05   %7d" % (openpyxl.load_workbook(
            os.path.join(target, files["register"]))["Group register"].max_row - 4),
        "WATCH_LIST_HISTORY           2026-10-01 06:07   %7d   cases opened from 2020-01-01" % rows("watchlist"),
        "FACILITY_OBLIGOR_EXTRACT     2026-10-01 01:15   %7d   both ledgers; held for the HB-DAC ledger close, ran "
        "after start of day 01/10" % rows("facilities"),
        "FACILITY_LIMIT_HISTORY       2026-10-08 05:50   %7d   limit versions to 07/10/2026" % rows("limits"),
        "TREASURY_CLOSING_RATES       2026-09-30 16:20   %7d" % rows("fx"),
        "",
        "Control: limit history reconciled to the facility extract, no breaks.",
        "Next quarter-end pack: 31/12/2026.",
    ]
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write("\n".join(lines) + "\n")
