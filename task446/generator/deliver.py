"""Golden deliverables for task446: the caseload workbook and the committee chart.

    python3 task446/generator/deliver.py

Every figure comes from solve.golden(), which reads target/ only. The workbook stores each figure at the
precision the prompt asks for; totals are rounded from the unrounded sums. verify() reads both files back.
"""
import datetime as dt
import html
import io
import os
import re
import sys
import zipfile

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import solve  # noqa: E402
import docs  # noqa: E402

GOLD = os.path.join(HERE, "..", "golden")
XLSX = "watchlist_inflow_2027.xlsx"
CHART = "watchlist_inflow_2027.svg"
BANDS = ["Under £5m", "£5m and under £15m", "£15m and under £40m", "£40m and under £100m", "£100m and over"]
PREPARED = dt.datetime(2026, 10, 27, 16, 18, 0)
AUTHOR = "Jill Berry"
TITLE = "%d of the %d groups we monitor are expected on the watch list by 30 September 2027"


def r1(x):
    return round(x + 1e-12, 1)


def figures(t=solve.TARGET):
    G = solve.golden(t)
    fc = G["fc"]
    rows = []
    for b, r in zip(BANDS, G["bands"]):
        rows.append(dict(band=b, groups=r["groups"], committed=r1(r["committed"] / 1e6), flagged=r["flagged"],
                         arrivals=r1(r["exp_arrivals"]), exposure=r1(r["exp_exposure"] / 1e6)))
    total = dict(band="Total", groups=sum(r["groups"] for r in G["bands"]),
                 committed=r1(sum(r["committed"] for r in G["bands"]) / 1e6),
                 flagged=sum(r["flagged"] for r in G["bands"]),
                 arrivals=r1(sum(r["exp_arrivals"] for r in G["bands"])),
                 exposure=r1(sum(r["exp_exposure"] for r in G["bands"]) / 1e6))
    cohorts = []
    for c, cc in zip(G["tab"][:4], fc["cc"]):
        cohorts.append(dict(date=dt.date.fromisoformat(c["date"]), **cc))
    F = dict(G=G, call=G["call"], rows=rows, total=total, cohorts=cohorts, measures=fc["measures"], N1=fc["N1"],
             F1=fc["F1"], D1=fc["D1"], I1=fc["I1"], a1=fc["a1"], p_d=fc["p_d"], p_i=fc["p_i"],
             p_unflag=fc["p_unflag"], expected=fc["total"], test=fc["test"],
             flagged_exp=fc["flagged"], unflagged_exp=fc["unflagged"])
    F["single"] = fc["test"]["hit_and_miss"]
    F["split"] = fc["test"]["split"]
    return F


# ------------------------------------------------------------------ workbook
GREY = "595F66"
HEAD_FILL = PatternFill("solid", fgColor="E9EBEE")
TOTAL_FILL = PatternFill("solid", fgColor="F4F5F7")
THIN = Side(style="thin", color="A6ABB2")


PREP_CREATED = dt.datetime(2026, 10, 21, 9, 37, 0)


def workbook(F, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Caseload 2027"
    ws["A1"] = "Halstow Bank plc  |  Credit Risk Monitoring"
    ws["A1"].font = Font(size=9, color=GREY)
    ws["A2"] = "Watch-list caseload for the year to 30 September 2027"
    ws["A2"].font = Font(bold=True, size=14)
    ws["A3"] = "Executive Risk Committee, 4 November 2026, item 3.1"
    ws["A3"].font = Font(size=10, color=GREY)
    ws["A5"] = "Groups expected to reach the watch list by 30 September 2027"
    ws["A5"].font = Font(bold=True, size=11)
    ws["F5"] = F["call"]
    ws["F5"].font = Font(bold=True, size=18)
    ws["F5"].number_format = "0"
    ws["F5"].alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[5].height = 27
    ws["A6"] = ("From the %d groups in the monitored book at 30 September 2026, %d of them flagged by the "
                "September early-warning run" % (F["N1"], F["F1"]))
    ws["A6"].font = Font(size=9, color=GREY)

    hdr_row = 8
    hdr = ["Caseload band\n(committed exposure at 30 Sep 2026)", "Groups", "Committed\nexposure (£m)",
           "Flagged in\nSeptember", "Expected\narrivals", "Expected committed exposure\narriving (£m)"]
    for c, h in enumerate(hdr, start=1):
        cell = ws.cell(row=hdr_row, column=c, value=h)
        cell.font = Font(bold=True, size=10)
        cell.fill = HEAD_FILL
        cell.border = Border(bottom=THIN)
        cell.alignment = Alignment(wrap_text=True, vertical="bottom", horizontal="left" if c == 1 else "right")
    ws.row_dimensions[hdr_row].height = 30
    r = hdr_row
    for row in F["rows"] + [F["total"]]:
        r += 1
        ws.cell(row=r, column=1, value=row["band"])
        ws.cell(row=r, column=2, value=row["groups"]).number_format = "#,##0"
        ws.cell(row=r, column=3, value=row["committed"]).number_format = "#,##0.0"
        ws.cell(row=r, column=4, value=row["flagged"]).number_format = "#,##0"
        ws.cell(row=r, column=5, value=row["arrivals"]).number_format = "0.0"
        ws.cell(row=r, column=6, value=row["exposure"]).number_format = "#,##0.0"
        if row["band"] == "Total":
            for c in range(1, 7):
                cell = ws.cell(row=r, column=c)
                cell.font = Font(bold=True)
                cell.fill = TOTAL_FILL
                cell.border = Border(top=THIN)
    last = r
    notes = [
        "1  Chance of reaching the watch list within the year, by the September flag: %.1f%% where the flag was "
        "raised on a customer we lend to, %.2f%% where it was raised only on a parent or sister company we do not lend "
        "to, %.2f%% where the group was not flagged. Set from the 2022 to 2025 cohorts taken together; see Cohorts."
        % (F["p_d"] * 100, F["p_i"] * 100, F["p_unflag"] * 100),
        "2  Committed exposure per CRM-07 s6.1: every facility on which a member is an obligor, once per group (a joint "
        "facility counts in each of its groups, and so does the facility of a company booked to one group and "
        "consolidated into another under s2.2), at its limit as it stood at the end of 30 September UK time (23:00 "
        "UTC in the limit history), in sterling at Group Treasury closing rates.",
        "3  Customers the grouping review had not yet booked are placed through their GLEIF consolidation (s2.2).",
        "4  Bands of £40m and over are worked by Restructuring, the rest by Business Support (s6.2).",
        "5  Computed unrounded. Caseload to the nearest whole group, band arrivals and exposure to one decimal (s5.3); "
        "totals rounded from the unrounded sums.",
    ]
    for i, n in enumerate(notes):
        cell = ws.cell(row=last + 2 + i, column=1, value=n)
        cell.font = Font(size=9, color=GREY)
    for col, wd in zip("ABCDEF", [34, 9, 14, 12, 12, 26]):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)
    ws.print_area = "A1:F%d" % (last + 1 + len(notes))
    ws.print_title_rows = "%d:%d" % (hdr_row, hdr_row)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0

    cohort_sheet(wb, F)
    notes_sheet(wb)
    wb.save(path)
    docs.finish_xlsx(path, AUTHOR, AUTHOR, PREP_CREATED, PREPARED)


def _head(ws, r, hdr, height=30):
    for c, h in enumerate(hdr, start=1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = Font(bold=True, size=10)
        cell.fill = HEAD_FILL
        cell.border = Border(bottom=THIN)
        cell.alignment = Alignment(wrap_text=True, vertical="bottom", horizontal="left" if c == 1 else "right")
    ws.row_dimensions[r].height = height


def cohort_sheet(wb, F):
    ws = wb.create_sheet("Cohorts")
    ws["A1"] = "Closed cohorts, by where the September flag was raised"
    ws["A1"].font = Font(bold=True, size=12)
    _head(ws, 3, ["Cohort", "Groups\nmonitored", "Flagged on a\ncustomer", "Reached", "Flagged only on\na "
                  "non-customer", "Reached", "Not\nflagged", "Reached", "Reached,\nall groups"])
    r = 3
    for c in F["cohorts"]:
        r += 1
        vals = [c["date"], c["N"], c["D"], c["D_R"], c["I"], c["I_R"], c["U"], c["U_R"], c["P"]]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=r, column=j, value=v)
        ws.cell(row=r, column=1).number_format = "d mmm yyyy"
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="left")
    m = F["measures"]
    r += 1
    tot = ["Four cohorts together", m["N"], m["D"], m["D_R"], m["I"], m["I_R"], m["U"], m["U_R"], m["P"]]
    for j, v in enumerate(tot, start=1):
        cell = ws.cell(row=r, column=j, value=v)
        cell.font = Font(bold=True)
        cell.fill = TOTAL_FILL
        cell.border = Border(top=THIN)
    r += 1
    rates = ["Share reaching the list", None, None, m["D_R"] / m["D"], None, m["I_R"] / m["I"], None,
             m["U_R"] / m["U"], m["P"] / m["N"]]
    for j, v in enumerate(rates, start=1):
        cell = ws.cell(row=r, column=j, value=v)
        if j > 1 and v is not None:
            cell.number_format = "0.0%"
    for rr in range(4, r):
        for j in range(2, 10):
            ws.cell(row=rr, column=j).number_format = "#,##0"

    r += 2
    ws.cell(row=r, column=1, value="Each closed cohort forecast from the other three taken together "
            "(forecast less actual arrivals, as a share of actual)").font = Font(bold=True, size=11)
    r += 1
    _head(ws, r, ["Basis", "2022", "2023", "2024", "2025", "Largest\nmiss", "2026 book,\nexpected"])
    for label, (miss, err, fw) in (("One rate for flagged groups, one for the rest", F["single"]),
                                   ("A rate by where the flag was raised, one for the rest", F["split"])):
        r += 1
        vals = [label] + list(miss) + [err, fw]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=j, value=v)
            if j in (2, 3, 4, 5):
                cell.number_format = "+0.0%;-0.0%;0.0%"
            elif j == 6:
                cell.number_format = "0.0%"
            elif j == 7:
                cell.number_format = "0.0"

    r += 2
    ws.cell(row=r, column=1, value="Book at 30 September 2026").font = Font(bold=True, size=11)
    lines = [
        ("Groups monitored", F["N1"], "#,##0", "active on the group register and not on the watch list (s4.1)"),
        ("Flagged AMBER or RED", F["F1"], "#,##0", "September 2026 publication, any member (s3.2)"),
        ("  raised on a customer", F["D1"], "#,##0", "%.1f%% expected to reach the list" % (F["p_d"] * 100)),
        ("  raised only on a non-customer", F["I1"], "#,##0", "%.2f%% expected to reach the list" % (F["p_i"] * 100)),
        ("Not flagged", F["N1"] - F["F1"], "#,##0", "%.2f%% expected to reach the list" % (F["p_unflag"] * 100)),
        ("Expected arrivals", F["expected"], "0.0", "%.1f from flagged groups, %.1f from the rest; caseload %d"
         % (F["flagged_exp"], F["unflagged_exp"], F["call"])),
    ]
    for label, v, fmt, note in lines:
        r += 1
        ws.cell(row=r, column=1, value=label)
        cell = ws.cell(row=r, column=2, value=v)
        cell.number_format = fmt
        ws.cell(row=r, column=3, value=note).font = Font(size=9, color=GREY)
    shares = [c["D"] / c["F"] for c in F["cohorts"]]
    r += 2
    for t in ("Customers raised %.0f%% to %.0f%% of the flags in the closed cohorts and %.0f%% of this September's."
              % (min(shares) * 100, max(shares) * 100, F["D1"] / F["F1"] * 100),):
        ws.cell(row=r, column=1, value=t).font = Font(size=9, color=GREY)
        r += 1
    for col, wd in zip("ABCDEFGHI", [46, 11, 12, 9, 14, 9, 10, 9, 11]):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A4"


def notes_sheet(wb):
    ws = wb.create_sheet("Notes")
    rows = [
        ("Notes", ""),
        ("As at", "Monitored book, flags and committed exposure at 30 September 2026. Forecast window 1 October 2026 "
                  "to 30 September 2027."),
        ("Obligor groups", "Group register customers plus GLEIF accounting-consolidation members in force at each date "
                           "(CRM-07 s2.2). Engine signals keyed by LEI or Companies House number."),
        ("Flag", "Any member AMBER or RED in the engine's month-end publication for the date (s3.2). A flag is raised "
                 "on a customer where the AMBER or RED entity is a customer on the credit system at the date."),
        ("Cohorts", "30 September 2022 to 2025. Reached: placed on the watch list, any watch level, within twelve "
                    "months of the cohort date; groups leaving the book unplaced have not reached it (s4.1, s4.2)."),
        ("Forecast", "Monitored book at 30 September 2026 only, no new groups (s5.1, s5.2). Each group's chance of "
                     "reaching the list is the 2022 to 2025 share for groups flagged the same way: flag raised on a "
                     "customer, flag raised only on a parent or sister company we do not lend to, or not flagged."),
        ("Committed exposure", "Facility and obligor extract run at 01:15 on 1 October, after the start-of-day batch; "
                               "limits taken from the limit history as at 23:00 UTC on 30 September, the end of the "
                               "day in BST (s1.3, s6.1). Facilities keyed by booking entity and facility number, "
                               "each counted once in each group with an obligor on it. Non-sterling limits divided by "
                               "the Group Treasury closing rate (units per pound)."),
        ("Bands", "Under £5m; £5m and under £15m; £15m and under £40m; £40m and under £100m; £100m and over (s6.2)."),
        ("Source", "Reference Data quarter-end pack as at 30/09/2026, compiled 08/10/2026."),
        ("Prepared", "Credit Risk Monitoring, 27 October 2026."),
    ]
    for a, b in rows:
        ws.append([a, b])
    ws["A1"].font = Font(bold=True, size=12)
    for rr in range(2, ws.max_row + 1):
        ws.cell(row=rr, column=1).font = Font(bold=True, size=10)
        ws.cell(row=rr, column=1).alignment = Alignment(vertical="top")
        ws.cell(row=rr, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 96


# ------------------------------------------------------------------ chart
def chart(F, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import matplotlib.transforms as mtransforms
    plt.rcParams["font.family"] = "Carlito"
    plt.rcParams["svg.fonttype"] = "none"
    plt.rcParams["svg.hashsalt"] = "halstow-crm-caseload-2027"
    ink, ink2, grid = "#1f2328", "#5b6169", "#e4e6ea"
    c_bar = "#a3412b"
    ar = [row["arrivals"] for row in F["rows"]]
    x = list(range(5))
    fig, ax = plt.subplots(figsize=(10, 6.1), dpi=150)
    fig.patch.set_facecolor("white")
    ax.bar(x, ar, width=0.58, color=c_bar, zorder=2)
    for i in x:
        ax.text(i, ar[i] + 0.5, "%.1f" % ar[i], ha="center", va="bottom", color=ink, fontsize=11, fontweight="bold",
                zorder=3)
    top = 10 * (int(max(ar) // 10) + 2)
    ax.set_xlim(-0.6, 4.6)
    ax.set_ylim(0, top)
    ax.set_yticks(list(range(0, top + 1, 10)))
    ax.set_yticklabels([str(v) for v in range(0, top + 1, 10)], fontsize=10, color=ink2)
    ax.set_ylabel("Expected arrivals (groups)", fontsize=10.5, color=ink2)
    ax.set_xticks(x)
    ax.set_xticklabels(["Under\n£5m", "£5m and\nunder £15m", "£15m and\nunder £40m", "£40m and\nunder £100m",
                        "£100m\nand over"], fontsize=10.5, color=ink)
    ax.tick_params(axis="x", length=0, pad=6)
    ax.tick_params(axis="y", length=0)
    ax.grid(axis="y", color=grid, linewidth=0.8, zorder=0)
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.spines["bottom"].set_color("#9aa0a8")
    tr = mtransforms.blended_transform_factory(ax.transData, ax.transAxes)
    for lo, hi, name in ((-0.32, 2.32, "Business Support"), (2.68, 4.32, "Restructuring")):
        ax.plot([lo, hi], [-0.175, -0.175], transform=tr, color="#9aa0a8", lw=0.9, clip_on=False)
        ax.text((lo + hi) / 2, -0.195, name, transform=tr, ha="center", va="top", fontsize=10, color=ink2)
    ax.text(0.985, 0.93, "%d" % F["call"], transform=ax.transAxes, ha="right", va="top", fontsize=46,
            fontweight="bold", color=c_bar)
    ax.text(0.985, 0.785, "groups in the 2027 watch-list caseload", transform=ax.transAxes, ha="right", va="top",
            fontsize=12, color=ink)
    ax.text(0.985, 0.73, "%.1f expected arrivals, to the nearest group" % F["total"]["arrivals"],
            transform=ax.transAxes, ha="right", va="top", fontsize=9.5, color=ink2)
    fig.text(0.012, 0.968, TITLE % (F["call"], F["N1"]), ha="left", va="top", fontsize=15, fontweight="bold", color=ink)
    fig.text(0.012, 0.922, "Expected arrivals by committed-exposure band at 30 September 2026", ha="left", va="top",
             fontsize=11, color=ink2)
    fig.text(0.012, 0.018, "A group flagged on a customer we lend to reaches the list with probability %.0f%%, one "
             "flagged only on a parent or sister company %.2f%%, one not flagged %.2f%% (2022 to 2025 cohorts). "
             "Bands per CRM-07 s6.2." % (F["p_d"] * 100, F["p_i"] * 100, F["p_unflag"] * 100), ha="left",
             va="bottom", fontsize=8.5, color=ink2)
    fig.text(0.988, 0.018, "Credit Risk Monitoring, 27 October 2026", ha="right", va="bottom", fontsize=8.5,
             color=ink2)
    fig.subplots_adjust(left=0.075, right=0.985, top=0.82, bottom=0.2)
    buf = io.StringIO()
    fig.savefig(buf, format="svg", facecolor="white",
                metadata={"Date": None, "Creator": None, "Format": None, "Type": None})
    plt.close(fig)
    svg = buf.getvalue()
    # Calibri first (Carlito is its metric twin, so the layout holds); plain group ids
    svg = svg.replace("font-family: 'Carlito'", "font-family: Calibri, Carlito, sans-serif")
    svg = re.sub(r'id="matplotlib\.', 'id="', svg)
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(svg)


# ------------------------------------------------------------------ read back
def verify(F):
    wb = openpyxl.load_workbook(os.path.join(GOLD, XLSX))
    ws = wb["Caseload 2027"]
    assert ws["F5"].value == F["call"]
    got = []
    for row in ws.iter_rows(min_row=9, max_row=14, values_only=True):
        got.append(tuple(row[:6]))
    want = [(r["band"], r["groups"], r["committed"], r["flagged"], r["arrivals"], r["exposure"])
            for r in F["rows"] + [F["total"]]]
    assert got == want, (got, want)
    assert sum(r["groups"] for r in F["rows"]) == F["total"]["groups"] == F["N1"]
    assert sum(r["flagged"] for r in F["rows"]) == F["total"]["flagged"] == F["F1"]
    for k in ("committed", "arrivals", "exposure"):
        assert abs(sum(r[k] for r in F["rows"]) - F["total"][k]) < 1e-6, k
    assert round(F["total"]["arrivals"]) == F["call"]
    with zipfile.ZipFile(os.path.join(GOLD, XLSX)) as z:
        meta = b"".join(z.read(n) for n in z.namelist() if n.startswith("docProps/"))
    assert b"openpyxl" not in meta and b"2026-10-2" in meta
    svg = open(os.path.join(GOLD, CHART), encoding="utf-8").read()
    texts = [html.unescape(t) for t in re.findall(r"<text[^>]*>([^<]*)</text>", svg)]
    assert "matplotlib" not in svg.lower() and "carlito'" not in svg.lower()
    assert TITLE % (F["call"], F["N1"]) in texts, texts
    assert str(F["call"]) in texts
    for row in F["rows"]:
        assert "%.1f" % row["arrivals"] in texts, row
    assert all("\u2014" not in t and "\u2013" not in t for t in texts)
    for wbk in (wb,):
        for s in wbk.worksheets:
            for row in s.iter_rows(values_only=True):
                for v in row:
                    if isinstance(v, str):
                        assert "\u2014" not in v and "\u2013" not in v, v
    return want


def main():
    os.makedirs(GOLD, exist_ok=True)
    F = figures()
    workbook(F, os.path.join(GOLD, XLSX))
    chart(F, os.path.join(GOLD, CHART))
    for w in verify(F):
        print("%-22s %5d %10.1f %5d %6.1f %9.1f" % w)
    print("caseload", F["call"], "p_d %.4f p_i %.4f p_u %.5f" % (F["p_d"], F["p_i"], F["p_unflag"]))


if __name__ == "__main__":
    main()
