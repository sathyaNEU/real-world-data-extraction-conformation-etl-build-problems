"""The pack's documents: the board's peat monitoring protocol and the 2024 survey report (PDF), the unit register
(XLSX), the 1997 laboratory printout and the data dictionary and contents (TXT)."""
from __future__ import annotations

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                    # noqa: E402
from reportlab.lib.pagesizes import A4                                              # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                     # noqa: E402
from reportlab.lib.units import cm                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                            # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                        # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle   # noqa: E402

import common as K   # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("Sans", "LiberationSans-Regular"), ("Sans-B", "LiberationSans-Bold"), ("Sans-I", "LiberationSans-Italic"),
              ("Serif", "LiberationSerif-Regular"), ("Serif-B", "LiberationSerif-Bold")):
    pdfmetrics.registerFont(TTFont(nm, FONT_DIR + f + ".ttf"))
pdfmetrics.registerFontFamily("Sans", normal="Sans", bold="Sans-B", italic="Sans-I", boldItalic="Sans-B")
pdfmetrics.registerFontFamily("Serif", normal="Serif", bold="Serif-B", italic="Serif", boldItalic="Serif-B")
INK, GREY, RULE = colors.HexColor("#1b2430"), colors.HexColor("#56616c"), colors.HexColor("#b8c0c8")
WATER = colors.HexColor("#1f4e79")
EARTH = colors.HexColor("#4a3b22")


def styles(face="Sans", accent=WATER):
    b = face + "-B"
    body = ParagraphStyle("b", fontName=face, fontSize=9.8, leading=13.4, textColor=INK, spaceAfter=5)
    return dict(
        body=body,
        h1=ParagraphStyle("h1", fontName=b, fontSize=14.5, leading=18, textColor=accent, spaceAfter=2),
        h2=ParagraphStyle("h2", fontName=b, fontSize=10.6, leading=14, textColor=INK, spaceBefore=8, spaceAfter=2,
                          keepWithNext=1),
        meta=ParagraphStyle("m", fontName=face, fontSize=8.4, leading=11.2, textColor=GREY, spaceAfter=9),
        cell=ParagraphStyle("c", fontName="Sans", fontSize=7.9, leading=9.9, textColor=INK),
        cellb=ParagraphStyle("cb", fontName="Sans-B", fontSize=7.9, leading=9.9, textColor=INK),
        note=ParagraphStyle("n", fontName="Sans", fontSize=7.7, leading=9.9, textColor=GREY, spaceBefore=2,
                            spaceAfter=6),
        clause=ParagraphStyle("p", parent=body, leftIndent=28, firstLineIndent=-28))


def _frame(left, right):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(GREY)
        canvas.setFont("Sans", 7.4)
        canvas.drawString(2.0 * cm, A4[1] - 1.25 * cm, left)
        canvas.drawRightString(A4[0] - 2.0 * cm, A4[1] - 1.25 * cm, right)
        canvas.drawRightString(A4[0] - 2.0 * cm, 1.1 * cm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _table(S, rows, widths, right=()):
    data = [[Paragraph(str(c), S["cellb"] if i == 0 else S["cell"]) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=[w * cm for w in widths], repeatRows=1, hAlign="LEFT")
    st = [("LINEBELOW", (0, 0), (-1, 0), .6, INK), ("LINEBELOW", (0, -1), (-1, -1), .4, RULE),
          ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 1.6),
          ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6)]
    for c in right:
        st.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
    t.setStyle(TableStyle(st))
    return t


def _clause(S, num, text):
    return Paragraph("<b>%s</b>&nbsp;&nbsp;%s" % (num, text), S["clause"])


def _doc(path, title, author, subject, left, right):
    return SimpleDocTemplate(path, pagesize=A4, leftMargin=2.0 * cm, rightMargin=2.0 * cm, topMargin=2.0 * cm,
                             bottomMargin=1.9 * cm, title=title, author=author, subject=subject, creator=author,
                             producer=author), _frame(left, right)


# --------------------------------------------------------------------------------------------- the protocol
def write_protocol(path):
    S = styles("Serif", WATER)
    doc, fr = _doc(path, "Lageveen peat monitoring network: protocol (English edition)", K.BOARD,
                   "Peat monitoring protocol", K.BOARD, "Peat monitoring network: protocol, edition 2024")
    st = [Paragraph("Lageveen peat monitoring network", S["h1"]),
          Paragraph("Protocol, edition 2024 (English edition). Adopted by the executive board of the %s on 15 January "
                    "2024; replaces the edition of 1996. Secretary: %s." % (K.BOARD, K.PEOPLE["secretary"]), S["meta"]),
          Paragraph("1. Purpose", S["h2"]),
          _clause(S, "1.1", "The network measures how much carbon the peat of each of the board's peat polder units "
                            "loses. The results serve the board's water-level orders (peilbesluiten) and its report to "
                            "the province on emissions from drained peat."),
          Paragraph("2. The network", S["h2"]),
          _clause(S, "2.1", "Twelve peat units carry three monitoring points each. The points were staked in 1996 and "
                            "cored for the baseline in spring 1997."),
          _clause(S, "2.2", "The pegs had gone by 2023. The points were re-staked in autumn 2023 from their 1996 "
                            "coordinates and given the network's current point codes; the point register "
                            "(meetpunten_veenmonitoring.csv) lists each point's 1997 and current code and both "
                            "positions."),
          Paragraph("3. Coring", S["h2"]),
          _clause(S, "3.1", "1997: one core a point to 1.00 m below the surface with a fixed-volume peat sampler, cut "
                            "in the field in slices of 5 cm. 2024: one core a point to 1.30 m with a sampler of the "
                            "same type, cut in slices of 10 cm, the laboratory's current slice length."),
          _clause(S, "3.2", "A slice's depths are measured from the surface at the time of coring. The surface level "
                            "of each point is levelled on the day it is cored."),
          Paragraph("4. Laboratory", S["h2"]),
          _clause(S, "4.1", "Each slice is analysed for dry bulk density (dried at 105 &deg;C, on the slice's field "
                            "volume), loss on ignition at 550 &deg;C for four hours, and total carbon by dry "
                            "combustion; loss on ignition and total carbon are expressed on the dry matter."),
          _clause(S, "4.2", "1997: %s, report 97-0412. 2024: %s, LIMS project \"HHLV veenmonitoring 2024\". Both "
                            "laboratories held accreditation for these methods and took part in the national soil "
                            "laboratories' proficiency scheme in the year of analysis." % (K.LAB97, K.LAB24)),
          Paragraph("5. The carbon loss", S["h2"]),
          _clause(S, "5.1", "A monitoring point's carbon loss is the carbon that the peat which formed the top 1.00 m "
                            "of the point's 1997 core has lost between the point's 1997 coring and its 2024 coring, "
                            "divided by the years between the two coring dates (the days between them divided by "
                            "365.25), in tonnes of carbon per hectare per year."),
          _clause(S, "5.2", "A unit's carbon loss is the mean of its three monitoring points' losses, reported to one "
                            "decimal."),
          Paragraph("6. Use of the results", S["h2"]),
          _clause(S, "6.1", "The water-level order of 2027 raises the summer target level first in the three units "
                            "with the largest carbon loss (5.2). The other units follow in the order's later phases."),
          _clause(S, "6.2", "The order states for every unit its carbon loss (5.2), its surface lowering since 1997 "
                            "(annex A) and its mean lowest groundwater level (annex B)."),
          Paragraph("Annex A. Surface lowering", S["h2"]),
          _clause(S, "A.1", "Each unit has a grid of levelling points, levelled in spring 1997 by the board's survey "
                            "team from the board's benchmarks and surveyed in spring 2024 by RTK-GNSS (survey report "
                            "of %s, 2024)." % K.SURVEY24),
          _clause(S, "A.2", "A unit's surface lowering is the mean, over its grid points, of the fall of the surface "
                            "at the point between the 1997 and the 2024 survey, divided by the years between the "
                            "point's two survey dates (days divided by 365.25), in millimetres per year, with the 1997 "
                            "levels brought to current NAP as the survey report states."),
          Paragraph("Annex B. Groundwater", S["h2"]),
          _clause(S, "B.1", "Two dipwells a unit, installed in November 2013, are read by pressure loggers at 08:00 "
                            "every day and by hand on the 14th and the 28th of every month. Both record the depth of "
                            "the water below the top of the tube."),
          _clause(S, "B.2", "A dipwell's mean lowest groundwater level (GLG) is the mean, over the hydrological years "
                            "(1 April to 31 March) 2017/18 to 2024/25, of each year's three deepest water levels among "
                            "the readings of the 14th and the 28th, in centimetres below the surface. The logger's "
                            "08:00 value of that day is the reading; where the logger has none, the hand reading of "
                            "that day is. A unit's GLG is the mean of its two dipwells."),
          Spacer(1, 6),
          Paragraph("The point register, the laboratory reports, the survey files and the dipwell series are kept by "
                    "the board's department of water systems.", S["note"])]
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ------------------------------------------------------------------------------------------ the survey report
def write_survey(path, n_grid, n_lost, n_units):
    S = styles("Sans", EARTH)
    doc, fr = _doc(path, "Surface levels Lageveen 2024: survey report", K.SURVEY24, "Surface level survey",
                   K.SURVEY24, "Surface levels Lageveen 2024")
    st = [Paragraph("Surface levels Lageveen 2024", S["h1"]),
          Paragraph("Survey report for the %s. %s, project 24-117. Surveyor: %s. 12 June 2024."
                    % (K.BOARD, K.SURVEY24, K.PEOPLE["surveyor"]), S["meta"]),
          Paragraph("Scope", S["h2"]),
          Paragraph("The levelling grids of the board's %d peat units (%d grid points) and the 36 points of the peat "
                    "monitoring network, which were re-staked by us in autumn 2023 and levelled on the day each was "
                    "cored in spring 2024." % (n_units, n_grid), S["body"]),
          Paragraph("Method", S["h2"]),
          Paragraph("RTK-GNSS with corrections from the national reference network, a five-second observation on "
                    "each point, heights in NAP through the current geoid model. Repeat observations on 40 points "
                    "agree to 0.011 m (standard deviation) in height and 0.02 m in position. Grid pegs were staked "
                    "out from the 1996 coordinates the board supplied. %d grid points could not be surveyed."
                    % n_lost, S["body"]),
          Paragraph("The 1997 levels", S["h2"]),
          Paragraph("The board's survey team levelled the grids and the monitoring points in 1997 from the board's "
                    "benchmarks, on the benchmark heights published in 1995. Those heights have since been "
                    "re-adjusted (2005 and 2018). To bring a 1997 level to current NAP, subtract %.3f m, the mean "
                    "re-adjustment of the eleven benchmarks the 1997 levelling started from (range 0.047 to "
                    "0.061 m)." % K.DATUM_OFFSET_M, S["body"]),
          Paragraph("Files delivered", S["h2"]),
          Paragraph("waterpassing_raster_1997_2024.csv: the grid points, one row a point and survey, the 1997 rows "
                    "transcribed from the board's levelling books. The monitoring points' levels are in the board's "
                    "point register.", S["body"])]
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ---------------------------------------------------------------------------------------- the unit register
def _style_sheet(sh, widths, fill="1F4E79"):
    from openpyxl.styles import Alignment, Font, PatternFill
    for i, w in enumerate(widths):
        sh.column_dimensions[chr(65 + i)].width = w
    for c in sh[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    sh.freeze_panes = "A2"


def write_register(path, reg):
    from openpyxl import Workbook
    wb = Workbook()
    sh = wb.active
    sh.title = "Peilvakken"
    sh.append(["unit_code", "unit_name", "area_ha", "land_use", "dominant_soil", "pumping_station",
               "summer_level_m_nap", "winter_level_m_nap"])
    for u in reg["units"]:
        sh.append([u["code"], u["name"], u["area"], u["landuse"], u["soil"], u["station"], u["summer"], u["winter"]])
    for row in sh.iter_rows(min_row=2):
        row[2].number_format = "0.0"
        row[6].number_format = "0.00"
        row[7].number_format = "0.00"
    _style_sheet(sh, [11, 13, 9, 36, 22, 22, 18, 18])
    ch = wb.create_sheet("Peilwijzigingen")
    ch.append(["unit_code", "effective", "season", "old_level_m_nap", "new_level_m_nap", "decision"])
    for c in reg["changes"]:
        ch.append([c["code"], c["date"], c["season"], c["old"], c["new"], c["reason"]])
    for row in ch.iter_rows(min_row=2):
        row[1].number_format = "dd-mm-yyyy"
        row[3].number_format = "0.00"
        row[4].number_format = "0.00"
    _style_sheet(ch, [11, 12, 9, 16, 16, 62])
    wb.save(path)


# ------------------------------------------------------------------------------------ text files
def write_lines(path, lines, encoding="utf-8"):
    with open(path, "w", encoding=encoding, newline="") as f:
        f.write("\r\n".join(lines) + "\r\n")


DICTIONARY = """HOOGHEEMRAADSCHAP VAN HET LAGEVEEN - department of water systems
Data dictionary for the peat monitoring and water system files (edition of 2 November 2026)

meetpunten_veenmonitoring.csv - the peat monitoring network's point register, one row a monitoring point
  unit_code, unit_name      the peat unit (peilvak) the point lies in
  point_1997, point_2024    the point's code in the 1997 baseline and its current code
  x_rd_1996, y_rd_1996      position staked in 1996 (RD New, m)
  x_rd_2023, y_rd_2023      position re-staked in autumn 2023 (RD New, m)
  cored_1997, cored_2024    coring dates
  surface_level_1997_m      surface level on the day of the 1997 coring, m, on the benchmark heights of 1995
  surface_level_2024_m      surface level on the day of the 2024 coring, m NAP
  core_length_1997_cm, core_length_2024_cm   depth the core reached below the surface

veenkernen_1997_bodemlab_wageveld.txt - Bodemlab Wageveld's report 97-0412 as printed, one line a slice
  Monsternr = sample number; Meetpunt = 1997 point code; Bemonsterd = sampling date (dd-mm-yyyy)
  Diepte (cm) = depth interval below the surface at coring; Vol.massa = dry bulk density, g/cm3
  Gloeiverl. = loss on ignition at 550 C, per cent of dry matter; C-totaal = total carbon, per cent of dry matter
  Numbers carry decimal commas. Page headers repeat on every page.

lims_export_veenkernen_2024.csv - Gouwlanden Laboratoria's LIMS export, one row a reported slice (semicolons, decimal commas)
  sample_id, project, point_code (current point code), sampling_date
  depth_top_cm, depth_bottom_cm   the slice's depth interval below the surface at coring
  field_description               the field crew's description of the slice (veen = peat, veraard = earthified,
                                  kleiig veen = clayey peat, venige klei = peaty clay, houtresten = wood remains)
  dry_bulk_density_kg_m3          dry bulk density, kg/m3
  loi_550_pct_dm, total_c_pct_dm, total_n_pct_dm   per cent of dry matter
  moisture_pct_fm                 water content, per cent of fresh mass
  analysis_date, report_no

lims_qc_2024.csv - the laboratory's quality control results for the project (semicolons, decimal commas)
  qc_type: duplicate (a second analysis of a reported sample) or reference material; recovery_pct against the
  reference value

waterpassing_raster_1997_2024.csv - the levelling grids, one row a grid point and survey
  unit_code, grid_point, campaign (1997 or 2024), survey_date, method, x_rd, y_rd (position of the shot, RD New, m)
  level_m (surface level, m), height_ref (1997: board benchmarks, heights of 1995; 2024: NAP)

peilbuizen_reeksen_2014_2026.csv - the dipwell network's readings, one row a reading
  dipwell_id, date_time (local time), source (logger: the pressure logger's 08:00 value; hand: the observer's
  reading), depth_below_tube_top_cm (depth of the water below the top of the tube), remark

peilbuizen_metadata.csv - one row a dipwell
  dipwell_id, unit_code, x_rd, y_rd, installed, tube_top_above_surface_m (height of the top of the tube above the
  surface), filter_bottom_m_below_surface, logger_serial

peilbuizen_onderhoud.csv - the dipwell maintenance log, one row a visit with work done
  date, dipwell_id, action, note, by (initials)

peilvakken_register_2026.xlsx - the board's register of peat units
  sheet Peilvakken: one row a unit, with its area, land use, dominant soil, pumping station and the summer and winter
  target levels in force (m NAP)
  sheet Peilwijzigingen: one row a change of a unit's target level, with the date it took effect

percelen_gewaspercelen_2025.csv - crop parcels declared for 2025 within the units, one row a parcel
  parcel_id, unit_code, area_ha, crop_code, crop, farm_id

bodemkaart_peilvakken.csv - the board's soil classes by unit, one row a unit and class, with its share of the unit's area

weerstation_gemaal_lageveen_2014_2026.csv - the board's weather station at the main pumping station, one row a day
  precipitation_mm, makkink_et_mm (reference evapotranspiration), t_mean_c
"""
