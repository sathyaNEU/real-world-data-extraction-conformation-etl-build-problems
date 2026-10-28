"""The goldens: segment_discards_2025.xlsx and segment_discards_2025.png, and generator/figures.json, every figure
read back from the pack."""
from __future__ import annotations

import json
import os

import analysis as A
import asks as AS
import common as K
import pack as PK


def half_up(x, nd=0):
    from decimal import ROUND_HALF_UP, Decimal
    q = Decimal(1).scaleb(-nd)
    return float(Decimal(repr(x)).quantize(q, rounding=ROUND_HALF_UP)) if nd else int(
        Decimal(repr(x)).quantize(Decimal(1), rounding=ROUND_HALF_UP))


def results(target):
    T, unmatched = PK.trip_records(target)
    shares = A.golden(T)
    fam = A.family(T)
    reg = PK.register(target)
    vessels = {s: sum(1 for r in reg.values() if r["segment"] == s) for s in K.SEG}
    parts = {}
    for s in K.SEG:
        X = [t for t in T if t["seg"] == s]
        parts[s] = dict(trips=len(X), observed=sum(t["observed"] for t in X))
    return dict(T=T, unmatched=unmatched, shares=shares, fam=fam, routes=A.routes(T), r2b=A.ratio_to_landings(T),
                bms=A.bms_rates(T), grounds=A.nursery_share(T), vessels=vessels, parts=parts,
                ask_a=AS.ask_a(target), ask_b=AS.ask_b(target), two=AS.two_terminal_hours(target))


def figures(res):
    sh = res["shares"]
    des = A.designated(sh)
    seg = {}
    for s in K.SEG:
        O = [t for t in res["T"] if t["seg"] == s and t["observed"]]
        N = [t for t in res["T"] if t["seg"] == s and not t["observed"]]
        seg[s] = dict(name=K.SEG_NAME[s], vessels=res["vessels"][s], trips=res["parts"][s]["trips"],
                      observed_trips=res["parts"][s]["observed"], share=half_up(sh[s]),
                      designated=s in des, first_sale_gbp_thousands=res["ask_a"][s], days_at_sea=res["ask_b"][s],
                      bms_kg_per_hour_observed=half_up(sum(t["land_bms"] for t in O) / sum(t["hours"] for t in O), 2),
                      bms_kg_per_hour_unobserved=half_up(sum(t["land_bms"] for t in N) / sum(t["hours"] for t in N),
                                                         2),
                      marketable_kg_per_hour_observed=half_up(sum(t["land_mk"] for t in O) /
                                                              sum(t["hours"] for t in O), 1),
                      marketable_kg_per_hour_unobserved=half_up(sum(t["land_mk"] for t in N) /
                                                                sum(t["hours"] for t in N), 1),
                      share_raised_by_marketable_landings=half_up(res["fam"]["ground/mk/decl"][s]))
    return {
        "threshold": K.THRESHOLD,
        "call": {"designated": list(des), "count": len(des),
                 "not_designated": [s for s in K.SEG if s not in des],
                 "nearest_over": min(des, key=lambda s: sh[s]),
                 "nearest_under": max((s for s in K.SEG if s not in des), key=lambda s: sh[s])},
        "segments": seg,
        "totals": {"vessels": sum(res["vessels"].values()),
                   "first_sale_gbp_thousands": half_up(sum(res["ask_a"].values()), 1),
                   "days_at_sea": sum(res["ask_b"].values()),
                   "trips": sum(p["trips"] for p in res["parts"].values()),
                   "observed_trips": sum(p["observed"] for p in res["parts"].values())},
        "default_path": {s: half_up(res["routes"]["R2 the observed discards per hour by ground, raised (default "
                                                  "path)"][s], 1) for s in K.SEG},
        "default_path_by_landed_weight_designated": list(A.designated(res["r2b"])),
        "vms_two_terminals": {"vessels": res["two"][0], "vessel_hours": res["two"][1], "apart_nm": K.APART_NM,
                              "apart_hours": res["two"][2]},
        "inshore_hours_share": {s: {"observed": half_up(100 * (res["grounds"][(s, True)]["WC14"] +
                                                                res["grounds"][(s, True)]["WC15"])),
                                    "unobserved": half_up(100 * (res["grounds"][(s, False)]["WC14"] +
                                                                  res["grounds"][(s, False)]["WC15"]))}
                                for s in K.SEG},
    }


def chart_png(path, fig):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    for f in ("LiberationSans-Regular.ttf", "LiberationSans-Bold.ttf"):
        p = os.path.join(K.FONT_DIR, f)
        if os.path.exists(p):
            font_manager.fontManager.addfont(p)
    plt.rcParams.update({"font.family": "Liberation Sans", "svg.hashsalt": "task275", "font.size": 10})
    seg = fig["segments"]
    order = sorted(K.SEG, key=lambda s: fig_share(fig, s))
    ink, muted, strong, grid = "#1d2733", "#5b6672", "#1f5a96", "#c9d1d9"
    fig_, ax = plt.subplots(figsize=(8.6, 4.9), dpi=150)
    ys = range(len(order))
    vals = [fig_share(fig, s) for s in order]
    cols = [strong if seg[s]["designated"] else "#b4bec8" for s in order]
    ax.barh(list(ys), vals, height=0.62, color=cols, edgecolor="white", linewidth=1.0, zorder=2)
    for y, s, v in zip(ys, order, vals):
        ax.text(v + 0.18, y, "%d%%" % seg[s]["share"], va="center", ha="left", fontsize=9.5, color=ink,
                fontweight="bold" if seg[s]["designated"] else "normal")
    ax.set_yticks(list(ys))
    ax.set_yticklabels(["%s  %s" % (s, short(s)) for s in order], fontsize=9.2, color=ink)
    ax.axvline(K.THRESHOLD, color=ink, linewidth=1.1, linestyle=(0, (4, 3)), zorder=3)
    ax.text(K.THRESHOLD + 0.12, len(order) - 0.45, "10% designation line (terms 2.1)", fontsize=8.6, color=ink,
            va="bottom", ha="left")
    ax.set_xlim(0, 16)
    ax.set_ylim(-0.6, len(order) - 0.1)
    ax.set_xticks(range(0, 17, 2))
    ax.set_xticklabels(["%d%%" % x for x in range(0, 17, 2)], fontsize=8.8, color=muted)
    ax.xaxis.grid(True, color=grid, linewidth=0.6, zorder=0)
    ax.set_axisbelow(True)
    for sp in ("top", "right", "left"):
        ax.spines[sp].set_visible(False)
    ax.spines["bottom"].set_color(grid)
    ax.tick_params(axis="y", length=0)
    ax.tick_params(axis="x", length=0)
    ax.set_xlabel("Discards as a share of catch, live weight, all 2025 trips", fontsize=9, color=muted)
    n = fig["call"]["count"]
    fig_.suptitle("%s of the eight fleet segments discarded more than 10%% of their 2025 catch" %
                  {4: "Four", 3: "Three", 5: "Five"}.get(n, str(n)), x=0.015, y=0.975, ha="left", fontsize=12.5,
                  fontweight="bold", color=ink)
    ax.set_title("Designated for the 2027 Selectivity Programme: %s" % ", ".join(fig["call"]["designated"]),
                 loc="left", fontsize=9.6, color=muted, pad=8)
    fig_.text(0.015, 0.018, "Source: WCFPO e-logbooks, landing declarations and At-Sea Observer Scheme, 2025. "
                            "Unobserved trips' discards raised by ground from the observed trips'\nunwanted catch "
                            "(BMS and lawful discards) per hour fished, less the BMS the unobserved trips declared.",
              fontsize=7.4, color=muted, ha="left", linespacing=1.4)
    fig_.subplots_adjust(left=0.30, right=0.97, top=0.84, bottom=0.19)
    fig_.savefig(path, dpi=150, metadata={"Software": None})
    plt.close(fig_)


SHORT = {"TR1": "Otter trawl 18 m+", "TR2": "Otter trawl 12-18 m", "NEP": "Nephrops trawl", "BT": "Beam trawl",
         "SN": "Seine", "GN": "Set gillnet", "LL": "Set longline", "POT": "Pots and creels"}


def short(s):
    return SHORT[s]


_UNROUNDED = {}


def fig_share(fig, s):
    return _UNROUNDED.get(s, fig["segments"][s]["share"])


def workbook(path, fig, res):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Segments 2025"
    navy = "173B5C"
    thin = Side(style="thin", color="B8C0C8")
    ws.append(["Wrayburn Coast FPO: fleet segments' 2025 discard shares, for the board of 8 December 2026"])
    ws["A1"].font = Font(bold=True, size=13, color=navy)
    ws.append(["Prepared by %s, science and compliance manager. Terms of the 2027 Selectivity Programme, clauses "
               "2.1 and 3.1 to 3.4." % K.REQUESTER])
    ws["A2"].font = Font(italic=True, size=9, color="56616C")
    ws.append([])
    head = ["Segment", "Description", "Vessels", "Trips 2025", "Observed trips", "Discard share 2025", "Designated",
            "First-sale value 2025 (GBP thousand)", "Days at sea 2025"]
    ws.append(head)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=navy)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[4].height = 46
    seg = fig["segments"]
    for s in K.SEG:
        d = seg[s]
        ws.append([s, d["name"], d["vessels"], d["trips"], d["observed_trips"], res["shares"][s] / 100.0,
                   "Yes" if d["designated"] else "No", d["first_sale_gbp_thousands"], d["days_at_sea"]])
        row = ws[ws.max_row]
        row[5].number_format = "0%"
        row[7].number_format = "#,##0.0"
        row[8].number_format = "#,##0"
        for c in (row[2], row[3], row[4]):
            c.number_format = "#,##0"
        if d["designated"]:
            for c in row:
                c.fill = PatternFill("solid", fgColor="E3ECF5")
            row[6].font = Font(bold=True, color=navy)
    tot = fig["totals"]
    ws.append(["All segments", None, tot["vessels"], tot["trips"], tot["observed_trips"], None, None,
               tot["first_sale_gbp_thousands"], tot["days_at_sea"]])
    row = ws[ws.max_row]
    for c in row:
        c.font = Font(bold=True)
        c.border = Border(top=thin)
    row[7].number_format = "#,##0.0"
    for c in (row[2], row[3], row[4], row[8]):
        c.number_format = "#,##0"
    ws.append([])
    ws.append(["Discard share: discards over catch, live weight, every trip begun in 2025 (terms 3.1 to 3.4). "
               "Designated: more than 10 per cent (terms 2.1)."])
    ws.append(["First-sale value: sales notes for landings of 2025 (terms 4.3). Days at sea: UK calendar days with a "
               "VMS report outside port limits (terms 4.2)."])
    for r in (ws.max_row - 1, ws.max_row):
        ws.cell(r, 1).font = Font(size=9, color="56616C")
    widths = {"A": 13, "B": 40, "C": 9, "D": 11, "E": 11, "F": 12, "G": 11, "H": 16, "I": 12}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    for r in ws.iter_rows(min_row=5, max_row=13):
        for c in r[2:]:
            c.alignment = Alignment(horizontal="right" if c.column != 7 else "center")
    ws.freeze_panes = "C5"
    ws.print_area = "A1:I16"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    w2 = wb.create_sheet("Observer coverage")
    w2.append(["Observed and unobserved trips, 2025: BMS declared on landing and the hours fished on the inshore "
               "grounds"])
    w2["A1"].font = Font(bold=True, size=12, color=navy)
    w2.append([])
    head2 = ["Segment", "Observed trips", "Unobserved trips", "BMS declared, observed trips (kg an hour fished)",
             "BMS declared, unobserved trips (kg an hour fished)", "Hours on WC14 and WC15, observed trips",
             "Hours on WC14 and WC15, unobserved trips"]
    w2.append(head2)
    for c in w2[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=navy)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    w2.row_dimensions[3].height = 44
    for s in K.SEG:
        d = seg[s]
        go = res["grounds"][(s, True)]
        gu = res["grounds"][(s, False)]
        w2.append([s, d["observed_trips"], d["trips"] - d["observed_trips"], d["bms_kg_per_hour_observed"],
                   d["bms_kg_per_hour_unobserved"], go["WC14"] + go["WC15"], gu["WC14"] + gu["WC15"]])
        row = w2[w2.max_row]
        row[1].number_format = "#,##0"
        row[2].number_format = "#,##0"
        row[3].number_format = "0.00"
        row[4].number_format = "0.00"
        row[5].number_format = "0%"
        row[6].number_format = "0%"
    w2.append([])
    for line in ["An observed trip lands the undersized quota fish (BMS) that the observer weighs, and fishes less on the",
                 "inshore grounds; an unobserved trip declares a fraction of it, except in GN, where most vessels land",
                 "their BMS on every trip. The unobserved trips' discards are therefore their unwanted catch (BMS and",
                 "lawful discards) at the rate per hour the same segment's observed trips show on the same ground, less",
                 "the BMS they declared. Weights: observer haul records and landing declarations, live weight."]:
        w2.append([line])
        w2.cell(w2.max_row, 1).font = Font(size=9, color="56616C")
    for k, v in {"A": 11, "B": 11, "C": 12, "D": 18, "E": 18, "F": 16, "G": 16}.items():
        w2.column_dimensions[k].width = v
    w2.freeze_panes = "B4"
    w2.page_setup.orientation = "landscape"
    w2.page_setup.fitToWidth = 1
    w2.sheet_properties.pageSetUpPr.fitToPage = True
    w2.print_area = "A1:G%d" % w2.max_row
    wb.save(path)


def write_figures(path, fig):
    with open(path, "w") as f:
        json.dump(fig, f, indent=1, sort_keys=True)
        f.write("\n")


def write(target, golden):
    os.makedirs(golden, exist_ok=True)
    res = results(target)
    fig = figures(res)
    _UNROUNDED.clear()
    _UNROUNDED.update(res["shares"])
    workbook(os.path.join(golden, K.GOLDEN), fig, res)
    chart_png(os.path.join(golden, K.CHART), fig)
    return fig, res
