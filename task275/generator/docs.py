"""The pack's documents: the selectivity programme's terms, the observer scheme's manual and the PO's data dictionary
(PDF); the quota uptake report (XLSX); and the pack's contents (TXT)."""
from __future__ import annotations

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                    # noqa: E402
from reportlab.lib.pagesizes import A4                                              # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                     # noqa: E402
from reportlab.lib.units import cm                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                            # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                        # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle   # noqa: E402

import common as K   # noqa: E402

FONT_DIR = K.FONT_DIR
for nm, f in (("Sans", "LiberationSans-Regular"), ("Sans-B", "LiberationSans-Bold"), ("Sans-I", "LiberationSans-Italic"),
              ("Serif", "LiberationSerif-Regular"), ("Serif-B", "LiberationSerif-Bold")):
    pdfmetrics.registerFont(TTFont(nm, FONT_DIR + f + ".ttf"))
pdfmetrics.registerFontFamily("Sans", normal="Sans", bold="Sans-B", italic="Sans-I", boldItalic="Sans-B")
pdfmetrics.registerFontFamily("Serif", normal="Serif", bold="Serif-B", italic="Serif", boldItalic="Serif-B")
INK, GREY, RULE = colors.HexColor("#1b2430"), colors.HexColor("#56616c"), colors.HexColor("#b8c0c8")
NAVY = colors.HexColor("#173b5c")
TEAL = colors.HexColor("#1d5a57")


def styles(face="Sans", accent=NAVY):
    b = face + "-B"
    body = ParagraphStyle("b", fontName=face, fontSize=9.8, leading=13.4, textColor=INK, spaceAfter=5)
    return dict(
        body=body,
        h1=ParagraphStyle("h1", fontName=b, fontSize=14.5, leading=18, textColor=accent, spaceAfter=2),
        h2=ParagraphStyle("h2", fontName=b, fontSize=10.6, leading=14, textColor=INK, spaceBefore=8, spaceAfter=2,
                          keepWithNext=1),
        meta=ParagraphStyle("m", fontName=face, fontSize=8.4, leading=11.2, textColor=GREY, spaceAfter=9),
        cell=ParagraphStyle("c", fontName="Sans", fontSize=7.9, leading=9.9, textColor=INK),
        cellb=ParagraphStyle("cb", fontName="Sans-B", fontSize=7.9, leading=9.9, textColor=INK),
        note=ParagraphStyle("n", fontName="Sans", fontSize=7.7, leading=9.9, textColor=GREY, spaceBefore=2,
                            spaceAfter=6),
        clause=ParagraphStyle("p", parent=body, leftIndent=28, firstLineIndent=-28))


def _frame(left, right):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(GREY)
        canvas.setFont("Sans", 7.4)
        canvas.drawString(2.0 * cm, A4[1] - 1.25 * cm, left)
        canvas.drawRightString(A4[0] - 2.0 * cm, A4[1] - 1.25 * cm, right)
        canvas.drawRightString(A4[0] - 2.0 * cm, 1.1 * cm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _table(S, rows, widths, right=()):
    data = [[Paragraph(str(c), S["cellb"] if i == 0 else S["cell"]) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=[w * cm for w in widths], repeatRows=1, hAlign="LEFT")
    st = [("LINEBELOW", (0, 0), (-1, 0), .6, INK), ("LINEBELOW", (0, -1), (-1, -1), .4, RULE),
          ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 1.6),
          ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6)]
    for c in right:
        st.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
    t.setStyle(TableStyle(st))
    return t


def _clause(S, num, text):
    return Paragraph("<b>%s</b>&nbsp;&nbsp;%s" % (num, text), S["clause"])


def _doc(path, title, author, subject, left, right):
    return SimpleDocTemplate(path, pagesize=A4, leftMargin=2.0 * cm, rightMargin=2.0 * cm, topMargin=2.0 * cm,
                             bottomMargin=1.9 * cm, title=title, author=author, subject=subject, creator=author,
                             producer=author), _frame(left, right)


# ------------------------------------------------------------------------------------------------ programme terms
def write_terms(path):
    S = styles("Serif", NAVY)
    doc, fr = _doc(path, "2027 Selectivity Programme: terms", K.PO, "Selectivity programme",
                   K.PO, "2027 Selectivity Programme: terms")
    st = [Paragraph("2027 Selectivity Programme", S["h1"]),
          Paragraph("Terms adopted by the board of the %s on 14 October 2026. Secretary to the board: %s, chief "
                    "executive." % (K.PO, K.PEOPLE["ceo"]), S["meta"]),
          Paragraph("1. The programme", S["h2"]),
          _clause(S, "1.1", "The programme fits each vessel of a designated segment with the selective gear the "
                            "PO's gear group has approved for its gear (square-mesh panels, separator grids, "
                            "benthic release panels or escape gaps) and with a camera recording the catch handling "
                            "on deck, and pays for the recorded footage to be reviewed."),
          _clause(S, "1.2", "The programme runs from 1 April 2027 for three years."),
          Paragraph("2. Designated segments", S["h2"]),
          _clause(S, "2.1", "A fleet segment is designated when its member vessels discarded more than 10 per cent "
                            "of their 2025 catch by weight."),
          _clause(S, "2.2", "The board fixes the designated segments at its meeting of 8 December 2026, on the "
                            "analysis of the PO's science and compliance manager. The designations stand for the "
                            "three years of the programme."),
          Paragraph("3. Definitions", S["h2"]),
          _clause(S, "3.1", "Catch is everything a vessel brought aboard on its fishing trips: what it landed and "
                            "what it returned to the sea."),
          _clause(S, "3.2", "Discards are the part of the catch returned to the sea."),
          _clause(S, "3.3", "A segment's 2025 catch and discards are those of every fishing trip its member "
                            "vessels began in 2025. A vessel's segment is the segment the PO's vessel register "
                            "gives it."),
          _clause(S, "3.4", "Weights are live weights. A landed weight in another presentation is converted with "
                            "the conversion factor of the PO's species list."),
          Paragraph("4. Costs and funding", S["h2"]),
          _clause(S, "4.1", "Gear and cameras are paid for from the PO's selectivity fund."),
          _clause(S, "4.2", "Footage is reviewed at a fixed rate for each day at sea of a fitted vessel. A vessel's "
                            "days at sea in a year are the calendar days, in UK time, on which at least one of its "
                            "VMS position reports lies outside the port limits in the PO's port limits table."),
          _clause(S, "4.3", "The fund is replenished by a levy on the first-sale value of the landings of the "
                            "designated segments' vessels. The board sets the levy rate each year from the "
                            "segments' first-sale value in the latest full year: the value at which the year's "
                            "landings were first sold, as the registered buyers' sales notes record it."),
          Paragraph("5. Review", S["h2"]),
          _clause(S, "5.1", "The board reviews the designations after the 2028 season, on the same terms."),
          Spacer(1, 10),
          Paragraph("Segments of the PO's vessel register", S["h2"]),
          _table(S, [["Segment", "Description", "Gear"]] +
                 [[c, n, g] for c, n, g, _, _ in K.SEGMENTS], [2.2, 9.5, 2.0])]
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ------------------------------------------------------------------------------------------------ observer manual
def write_manual(path):
    S = styles("Sans", TEAL)
    doc, fr = _doc(path, "Observer scheme manual 2025", K.PEOPLE["coordinator"], "At-sea observer scheme",
                   K.SCHEME, "Scheme manual 2025")
    st = [Paragraph("%s: scheme manual 2025" % K.SCHEME, S["h1"]),
          Paragraph("Issued by %s, scheme coordinator, January 2025. Replaces the 2023 manual." %
                    K.PEOPLE["coordinator"], S["meta"]),
          Paragraph("1. What the scheme does", S["h2"]),
          Paragraph("The scheme places trained observers on member vessels to record, haul by haul, what is "
                    "caught, what is kept and what goes back to the sea. Its records feed the PO's catch "
                    "planning, the stock assessments the PO contributes to and the PO's reports to its members. "
                    "Observers are employed by the PO and work to this manual.", S["body"]),
          Paragraph("2. Selecting trips", S["h2"]),
          Paragraph("Skippers notify every fishing trip through the PO's trip app at least 24 hours before the "
                    "planned departure, giving the planned departure time. Each evening the coordinator's "
                    "selection program draws from the trips notified for the next day at random: about one trip "
                    "in twelve, and about one beam trawl trip in nine, the higher rate set at the request of the "
                    "sole and plaice management group. A drawn trip carries an observer when one is free that day; "
                    "when none is, the trip sails unobserved. A drawn trip that does not sail is recorded as "
                    "cancelled.", S["body"]),
          Paragraph("Every drawn trip is entered in the deployment log with the vessel and the planned departure "
                    "as the notification gives them, the outcome and, for an observed trip, the observer and the "
                    "observer's trip reference.", S["body"]),
          Paragraph("3. At sea", S["h2"]),
          Paragraph("The observer samples every haul (every set, for static gear) that they can attend; on a long "
                    "trip a haul shot during the observer's rest is not sampled, and the haul records hold the hauls "
                    "sampled. The observer takes no part in sorting the catch, and weighs each species in each of "
                    "three fractions, subsampling large hauls and raising the subsample to the haul:", S["body"]),
          _table(S, [["Fraction", "What it holds"],
                     ["RET", "catch retained to be landed for human consumption"],
                     ["BMS", "catch of a species under the landing obligation below its minimum conservation "
                             "reference size, retained to be landed separately"],
                     ["DIS", "catch returned to the sea, whatever the species and size"]], [2.2, 13.0]),
          Spacer(1, 6),
          Paragraph("Weights are live weights in kilograms, raised to the whole haul. Times are UK local times. "
                    "The position is where the gear was shot; the depth is the depth at the shot.", S["body"]),
          Paragraph("4. Checking", S["h2"]),
          Paragraph("The coordinator checks each trip's records within ten working days of its return: every "
                    "haul present, the fractions complete for every species weighed, the haul times in order. "
                    "Records that fail a check are returned to the observer and corrected before they are "
                    "loaded.", S["body"]),
          Paragraph("5. Contacts", S["h2"]),
          Paragraph("Scheme coordinator: %s. Data: %s, data officer." % (K.PEOPLE["coordinator"],
                                                                        K.PEOPLE["data"]), S["body"])]
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ------------------------------------------------------------------------------------------------ data dictionary
DICTIONARY = [
    ("vessel_register.csv", "One row a member vessel in 2025.", [
        ("rss_no", "Registry of Shipping and Seamen number"),
        ("vessel_pln", "port letters and number"),
        ("segment", "the PO's fleet segment (the programme terms list them)"),
        ("gear", "main gear: OTB otter trawl, TBB beam trawl, SDN seine, GNS set gillnet, LLS set longline, FPO "
                 "pots and creels"),
        ("overall_length_m, engine_kw, year_built, home_port, member_since", "as registered with the PO")]),
    ("elogbook_trips_2025.csv", "One row a fishing trip begun in 2025, from the skippers' electronic logbooks.", [
        ("trip_id", "the logbook's trip number"),
        ("departure_utc, return_utc", "departure from and return to port, UTC"),
        ("po_ground", "the PO ground the trip fished, from the logbook's fishing positions"),
        ("hauls, fishing_hours", "hauls (sets, for static gear) and the hours the gear was fishing"),
        ("discards_declared_kg", "discards the skipper declared in the logbook, which asks for discards of more "
                                 "than 50 kg of a species")]),
    ("elogbook_landings_2025.csv", "One row a species, presentation and size status of a trip's landing "
                                   "declaration, as the PO's catch database holds it.", [
        ("landing_date, landing_port", "where and when the catch was landed (UK date)"),
        ("mcrs_status", "Above MCRS, or BMS: below the minimum conservation reference size, landed separately "
                        "and not for human consumption"),
        ("presentation", "GUT gutted, WHL whole"),
        ("live_weight_kg", "the declared weight converted to live weight with the species list's conversion "
                           "factor")]),
    ("observer_deployments_2025.csv", "One row a trip drawn by the observer scheme's selection program.", [
        ("notified_departure", "planned departure in the skipper's notification, UK time"),
        ("vessel_pln, vessel_name", "as the notification gives them"),
        ("outcome", "Observed, No observer available, Trip cancelled"),
        ("observer_trip_ref", "keys the haul records of an observed trip"),
        ("hauls_sampled", "hauls the observer sampled; the haul records hold these hauls only")]),
    ("observer_hauls_2025.csv", "One row a species and fraction of an observed haul.", [
        ("shot_local, hauled_local", "UK time the gear was shot and hauled"),
        ("shot_lat, shot_lon, depth_m", "position and depth at the shot"),
        ("fraction", "RET retained, BMS retained below MCRS, DIS returned to the sea"),
        ("weight_kg", "live weight raised to the haul")]),
    ("sales_notes_2025.csv", "One row a line of a sales note submitted by a registered buyer for the first sale "
                             "of a member vessel's landing; notes with a sale date from 1 January 2025 to 31 January "
                             "2026 (extract of 6 February 2026).", [
        ("sales_note_no, buyer_reg_no, buyer_name", "the note and the registered buyer that submitted it"),
        ("sale_date, landing_date, port_of_landing", "dates are UK dates"),
        ("size_grade", "the buyer's size grade, 1 the largest"),
        ("weight_kg, price_gbp_per_kg, value_gbp", "product weight sold as the buyer weighed it, the price and "
                                                   "weight times price"),
        ("submitted_utc", "time the note reached the registered buyers' system")]),
    ("vms_positions_2025.csv", "One row a position report in 2025 from the VMS terminals on members' vessels, as "
                               "supplied by the PO's VMS service provider.", [
        ("terminal_id", "the terminal's serial number"),
        ("vessel_pln, vessel_name", "the vessel the terminal is assigned to"),
        ("position_utc", "time of the report, UTC"),
        ("latitude, longitude, speed_kn, course_deg", "as reported")]),
    ("vms_transponders.csv", "One row a member vessel's VMS terminal, as the PO's equipment record holds it at the end "
                             "of 2025.", [
        ("terminal_id, terminal_model", "the terminal fitted to the vessel"),
        ("previous_terminal_id", "the terminal the vessel carried before it; blank where the record holds none")]),
    ("port_limits.csv", "One row a port the members land at.", [
        ("harbour_lat, harbour_lon, port_limit_radius_nm", "the harbour reference point and the radius of the "
                                                           "port limits about it, in nautical miles"),
        ("fish_market", "Y where the port has a fish market selling by auction")]),
    ("fishing_grounds.csv", "One row a PO ground.", [
        ("lat_min to lon_max", "the ground's bounds")]),
    ("species_list.csv", "One row a species the members land or catch.", [
        ("landing_obligation", "Y for a species that must be landed under the landing obligation"),
        ("mcrs_cm", "minimum conservation reference size"),
        ("survivability_exemption", "Y where the exemption allows the species to be returned"),
        ("conversion_factor", "live weight over product weight in the usual presentation (1.00 whole)")]),
    ("quota_uptake_2025.xlsx", "The PO's quota uptake report for 2025: allocation, transfers and landings by stock "
                               "and month, in tonnes.", []),
    ("gear_specifications_2025.csv", "One row a gear declaration of a member vessel in 2025.", []),
]


def write_dictionary(path):
    S = styles("Sans", NAVY)
    doc, fr = _doc(path, "Data dictionary 2025", K.PEOPLE["data"], "Member data", K.PO, "Data dictionary 2025")
    st = [Paragraph("%s: data dictionary 2025" % K.PO, S["h1"]),
          Paragraph("%s, data officer. Revised 2 November 2026." % K.PEOPLE["data"], S["meta"])]
    for name, grain, cols in DICTIONARY:
        st.append(Paragraph(name, S["h2"]))
        st.append(Paragraph(grain, S["body"]))
        if cols:
            st.append(_table(S, [["Column", "Meaning"]] + [list(c) for c in cols], [5.2, 11.6]))
            st.append(Spacer(1, 4))
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ------------------------------------------------------------------------------------------------ quota uptake
def _style_sheet(sh, widths=None, fill="173B5C"):
    from openpyxl.styles import Alignment, Font, PatternFill
    for cell in sh[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sh.freeze_panes = "B2"
    for col in sh.columns:
        letter_ = col[0].column_letter
        if widths and letter_ in widths:
            sh.column_dimensions[letter_].width = widths[letter_]
        else:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:200])
            sh.column_dimensions[letter_].width = min(40, max(8, width + 2))


def write_quota(path, rows):
    """rows: (stock, allocation t, in t, out t, [12 monthly landings t])"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Uptake 2025"
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    ws.append(["Stock", "Initial allocation (t)", "Transfers in (t)", "Transfers out (t)", "Adjusted allocation (t)"]
              + months + ["Landed 2025 (t)", "Uptake"])
    for stock, alloc, tin, tout, mon in rows:
        adj = alloc + tin - tout
        tot = sum(mon)
        ws.append([stock, alloc, tin, tout, round(adj, 1)] + [round(m, 1) for m in mon] + [round(tot, 1),
                                                                                           round(tot / adj, 4)])
    for row in ws.iter_rows(min_row=2):
        for c in row[1:-1]:
            c.number_format = "#,##0.0"
        row[-1].number_format = "0.0%"
    _style_sheet(ws, {"A": 26})
    w2 = wb.create_sheet("Notes")
    for line in ["Quota uptake, %s, 2025 (closed 30 January 2026)." % K.PO,
                 "Landings are live weights from the members' landing declarations and include BMS landings, which "
                 "count against quota.",
                 "Transfers are swaps and leases with other producer organisations, as agreed with the "
                 "administration."]:
        w2.append([line])
    w2.column_dimensions["A"].width = 110
    wb.save(path)


def write_contents(path, items):
    out = ["%s: discard shares review pack" % K.PO,
           "Assembled %s by the science and compliance team." % K.ASSEMBLED.strftime("%d %B %Y").lstrip("0"), ""]
    for name, desc in items:
        out.append(name)
        out.append("    " + desc)
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(out) + "\n")
