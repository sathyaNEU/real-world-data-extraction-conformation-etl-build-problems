"""The golden deliverables, computed forward from the shipped files: rural_lender_recruitment_2027.xlsx (the five to
recruit first in order with counts, the sixth, every shortlisted lender, and the county figures for 2024 and 2025 and
for the July 2025 to June 2026 program year) and rural_lender_recruitment_2027.pptx (the twelve lenders' counts with
the five marked), with generator/figures.json for the write-up."""
import datetime as dt
import json
import os

import calc as C

XLSX = "rural_lender_recruitment_2027.xlsx"
PPTX = "rural_lender_recruitment_2027.pptx"
SHORT = {"Corvane Home Loans, LLC": "Corvane Home Loans", "Halverd Mortgage Company": "Halverd Mortgage",
         "Tessaline Home Lending, Inc.": "Tessaline Home Lending", "Arkhaven Bank, N.A.": "Arkhaven Bank",
         "Shunga Valley Bank": "Shunga Valley Bank", "Ninnescah Federal Credit Union": "Ninnescah FCU",
         "Saltfork State Bank": "Saltfork State Bank", "Weyrauch State Bank": "Weyrauch State Bank",
         "Mennwell Credit Union": "Mennwell Credit Union", "Heidebrecht State Bank": "Heidebrecht State Bank",
         "Kessinger Federal Credit Union": "Kessinger FCU", "Eisenbrandt State Bank": "Eisenbrandt State Bank"}
COUNTY_ROWS = ["Saline", "McPherson", "Dickinson", "Marion", "other four"]
COUNTY_LABEL = {"Saline": "Saline", "McPherson": "McPherson", "Dickinson": "Dickinson", "Marion": "Marion",
                "other four": "Rice, Ottawa, Ellsworth and Lincoln together"}


def fmt(n):
    return "{:,}".format(int(n))


def analyse(target):
    B = C.load(target)
    caps, route = C.captures(B)
    v = C.lender_counts(B, "golden")
    order = sorted(v.items(), key=lambda kv: (-kv[1], kv[0]))
    feed = C.counts(B, lambda F: F.mortgage)
    names = dict((name, n) for n, name in B["short"])
    reporter = {name: B["lei_of"].get(n, "") in B["hmda"].index for n, name in B["short"]}
    rows = {name: int(feed.get(names[name], 0)) for name in v}
    hm_route = {}
    fe_route = {}
    for n in B["reporter_nmls"]:
        r = route(n)
        hm_route[r] = hm_route.get(r, 0) + int(B["hmda"][B["lei_of"][n]])
        fe_route[r] = fe_route.get(r, 0) + int(feed.get(n, 0))
    pooled = C.capture(B)["all"]
    two, py = C.ask_figures(C.market_two_year(B), C.program_year(B))
    default = C.lender_counts(B, "unscaled")
    pooled_v = C.lender_counts(B, "pooled")
    r4 = C.estimate(B, sel=lambda F: F.mortgage | F.qcd)
    return dict(B=B, caps=caps, route=route, v=v, order=order, rows=rows, reporter=reporter, hm_route=hm_route,
                fe_route=fe_route, pooled=pooled, two=two, py=py, default=default, pooled_v=pooled_v, r4=r4,
                routes={name: route(names[name]) for name in v})


def figures(X):
    order = X["order"]
    five = [n for n, _ in order[:5]]
    sixth = order[5][0]
    d_top, d_six = C.five(X["default"])
    p_top, p_six = C.five(X["pooled_v"])
    r4_top, r4_six = C.five(X["r4"])
    f = dict(
        five=[SHORT[n] for n in five],
        five_counts={SHORT[n]: fmt(C.r10(X["v"][n])) for n in five},
        sixth=SHORT[sixth], sixth_count=fmt(C.r10(X["v"][sixth])),
        twelve={SHORT[n]: fmt(C.r10(x)) for n, x in order},
        reporter_counts={SHORT[n]: fmt(x) for n, x in order if X["reporter"][n]},
        capture_packaged="%d" % round(100 * X["caps"]["P"]), capture_separate="%d" % round(100 * X["caps"]["S"]),
        capture_pooled="%d" % round(100 * X["pooled"]),
        reporters_packaged_hmda=fmt(X["hm_route"]["P"]), reporters_packaged_rows=fmt(X["fe_route"]["P"]),
        reporters_separate_hmda=fmt(X["hm_route"]["S"]), reporters_separate_rows=fmt(X["fe_route"]["S"]),
        recorded_rows={SHORT[n]: fmt(X["rows"][n]) for n, _ in order if not X["reporter"][n]},
        two_year={k: fmt(C.r10(X["two"][k])) for k in COUNTY_ROWS},
        program_year={k: fmt(C.r10(X["py"][k])) for k in COUNTY_ROWS + ["total"]},
        default_five=[SHORT[n] for n, _ in d_top], default_sixth=SHORT[d_six[0]],
        pooled_five=[SHORT[n] for n, _ in p_top],
        pooled_counts={SHORT[n]: fmt(C.r10(x)) for n, x in p_top},
        weyrauch_unscaled=fmt(C.r10(X["default"]["Weyrauch State Bank"])),
        mennwell_pooled=fmt(C.r10(X["pooled_v"]["Mennwell Credit Union"])),
        weyrauch_pooled=fmt(C.r10(X["pooled_v"]["Weyrauch State Bank"])),
        quitclaims_kept_five=[SHORT[n] for n, _ in r4_top],
        heidebrecht_quitclaims_kept=fmt(C.r10(X["r4"]["Heidebrecht State Bank"])),
        heidebrecht_quitclaim_rows=fmt(int(((X["B"]["F"].qcd) & (X["B"]["F"].nmls == dict(
            (nm, n) for n, nm in X["B"]["short"])["Heidebrecht State Bank"]) &
            X["B"]["F"].MTG_INSTR_DATE.between(*C.WIN)).sum())),
        margin_fifth_sixth="%.2f" % (X["v"][five[4]] / X["v"][sixth]),
        clauses=["7.2", "7.3", "7.4", "7.5", "2.2"],
        nearest=10, years="2024 and 2025", program_year_span="1 July 2025 to 30 June 2026",
    )
    return f, five, sixth


def write_xlsx(path, X, f, five, sixth):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Recruitment order"
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="3B5B3A")
    mark = PatternFill("solid", fgColor="E3ECD9")
    thin = Side(style="thin", color="7F8F7E")
    ws.append(["2027 rural down-payment program: lenders to recruit first"])
    ws.append(["First-lien home-purchase mortgages on homes in the eight counties, closed in 2024 and 2025, "
               "nearest 10"])
    ws.append([])
    ws.append(["Rank", "Lender", "Mortgages, 2024 and 2025", "Recruit first", "Source of the count"])
    for k, (name, x) in enumerate(X["order"], 1):
        src = ("HMDA, reported originations" if X["reporter"][name] else
               "Recorder file, %s rows over the capture of mortgages recorded %s" % (
                   fmt(X["rows"][name]), "with the deed" if X["routes"][name] == "P" else "after the deed"))
        ws.append([k, name, C.r10(x), "Yes" if k <= 5 else ("Sixth" if k == 6 else None), src])
    for c in ws[4]:
        c.font = head
        c.fill = fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in range(5, 5 + len(X["order"])):
        ws.cell(r, 3).number_format = "#,##0"
        if r - 4 <= 5:
            for col in range(1, 6):
                ws.cell(r, col).fill = mark
        if r - 4 == 5:
            for col in range(1, 6):
                ws.cell(r, col).border = Border(bottom=thin)
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"].font = Font(italic=True, size=9, color="5F5F5A")
    for col, wd in zip("ABCDE", (6, 34, 14, 12, 62)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A5"
    ws2 = wb.create_sheet("County volumes")
    ws2.append(["First-lien home-purchase mortgages on homes in each county, every lender, nearest 10"])
    ws2.append([])
    ws2.append(["County", "Closed in 2024 and 2025", "Closed 1 July 2025 to 30 June 2026"])
    for k in COUNTY_ROWS:
        ws2.append([COUNTY_LABEL[k], C.r10(X["two"][k]), C.r10(X["py"][k])])
    ws2.append(["All eight counties, program year", None, C.r10(X["py"]["total"])])
    for c in ws2[3]:
        c.font = head
        c.fill = fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in range(4, 4 + len(COUNTY_ROWS) + 1):
        for col in (2, 3):
            ws2.cell(r, col).number_format = "#,##0"
    last = 4 + len(COUNTY_ROWS)
    for col in range(1, 4):
        ws2.cell(last, col).font = Font(bold=True)
        ws2.cell(last, col).border = Border(top=thin)
    ws2["A1"].font = Font(bold=True, size=12)
    for col, wd in zip("ABC", (44, 22, 30)):
        ws2.column_dimensions[col].width = wd
    ws2.freeze_panes = "A4"
    ws3 = wb.create_sheet("Notes")
    notes = [
        "Counts are first-lien mortgages a bank, credit union or licensed mortgage company made to finance an "
        "arm's-length purchase of a home in Saline, McPherson, Dickinson, Marion, Rice, Ottawa, Ellsworth or Lincoln "
        "County, dated by closing (program guidelines 7.1 to 7.5).",
        "Lenders that report HMDA are counted from the 2024 and 2025 loan-level file: originations (action 1), home "
        "purchase (purpose 1), first lien, property in the eight counties.",
        "Lenders that do not report HMDA are counted from the Venridge purchase-money file, after removing quitclaim "
        "refinances into a family trust, re-recordings, second mortgages against the same deed, seller-carried notes "
        "and the second county's record of a farmstead mortgage, then divided by the share of the HMDA reporters' "
        "mortgages the file holds for the way the mortgage was recorded: %s of the reporters' %s mortgages recorded the "
        "same day as the deed (%s per cent), %s of their %s recorded after it (%s per cent)." % (
            f["reporters_packaged_rows"], f["reporters_packaged_hmda"], f["capture_packaged"],
            f["reporters_separate_rows"], f["reporters_separate_hmda"], f["capture_separate"]),
        "County figures: a home counts in the county in which the property is recorded (program guidelines 2.2), so "
        "each HMDA reporter's count is spread over the counties its recorded mortgages stand in (HMDA places some "
        "rural homes in the county of their postal city); the program year is counted from the "
        "recorder file for every lender, each mortgage over the capture of the way it was recorded.",
        "Every count is estimated from unrounded figures and rounded to the nearest 10 once; the program year's total "
        "for all eight counties is rounded from the unrounded county figures, so it need not equal the sum of the "
        "rounded rows.",
        "Prepared by the research desk, 9 October 2026, from the files pulled for the December board.",
    ]
    ws3.append(["Notes"])
    for n in notes:
        ws3.append([n])
    ws3["A1"].font = Font(bold=True, size=12)
    ws3.column_dimensions["A"].width = 120
    for row in ws3.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def write_pptx(path, X, f, five, sixth):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    from pptx.util import Emu, Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Recruit %s, %s, %s, %s and %s first" % tuple(f["five"])
    p.font.size, p.font.bold = Pt(26), True
    p.font.color.rgb = RGBColor(0x2B, 0x44, 0x2A)
    p2 = tf.add_paragraph()
    p2.text = ("First-lien home-purchase mortgages on homes in the eight counties, closed in 2024 and 2025, nearest 10. "
               "Two of the five, Weyrauch State Bank and Mennwell Credit Union, do not report HMDA.")
    p2.font.size = Pt(14)
    p2.font.color.rgb = RGBColor(0x5F, 0x5F, 0x5A)
    order = X["order"]
    names = [SHORT[n] for n, _ in order][::-1]
    vals = [C.r10(x) for _, x in order][::-1]
    cd = CategoryChartData()
    cd.categories = names
    cd.add_series("Mortgages, 2024 and 2025", vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.95), Inches(12.1), Inches(4.8), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    plot = ch.plots[0]
    plot.gap_width = 55
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = "#,##0", False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(12)
    ser = plot.series[0]
    n = len(names)
    for i in range(n):
        pt = ser.points[i]
        pt.format.fill.solid()
        rank = n - i
        pt.format.fill.fore_color.rgb = RGBColor(0x3B, 0x5B, 0x3A) if rank <= 5 else RGBColor(0xC9, 0xD2, 0xC4)
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(12)
    va = ch.value_axis
    va.has_major_gridlines = False
    va.tick_labels.font.size = Pt(11)
    va.maximum_scale = 350
    va.minimum_scale = 0
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = "#,##0", False
    leg = slide.shapes.add_textbox(Inches(9.0), Inches(5.45), Inches(3.8), Inches(0.6))
    lp = leg.text_frame.paragraphs[0]
    lp.text = "Dark bars: the five recruited first"
    lp.font.size = Pt(12)
    lp.font.color.rgb = RGBColor(0x3B, 0x5B, 0x3A)
    src = slide.shapes.add_textbox(Inches(0.6), Inches(6.8), Inches(12.1), Inches(0.55))
    src.text_frame.word_wrap = True
    sp = src.text_frame.paragraphs[0]
    sp.text = ("Sources: HMDA loan-level file 2024 and 2025 (reporting lenders); Venridge purchase-money file, recordings "
               "January 2024 to August 2026, scaled by its capture of the reporters' mortgages (the others). "
               "Prairie Lantern Housing Fund research desk, October 2026.")
    sp.font.size = Pt(9)
    sp.font.color.rgb = RGBColor(0x5F, 0x5F, 0x5A)
    num = slide.shapes.add_textbox(Inches(12.55), Inches(7.0), Inches(0.5), Inches(0.3))
    np_ = num.text_frame.paragraphs[0]
    np_.text = "1"
    np_.font.size = Pt(9)
    np_.font.color.rgb = RGBColor(0x5F, 0x5F, 0x5A)
    prs.save(path)


def scrub_pptx(path, created, modified):
    """PowerPoint keeps a chart's data in an embedded workbook; the writing library stamps it with the build clock.
    Give it the deck's own stamps and Office's 1980 entry dates, and drop the library's stock thumbnail."""
    import io
    import re
    import zipfile
    with zipfile.ZipFile(path) as z:
        items = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    out = []
    for name, data in items:
        if name == "docProps/thumbnail.jpeg":
            continue
        if name == "_rels/.rels":
            data = re.sub(rb'<Relationship [^>]*Target="docProps/thumbnail.jpeg"/>', b"", data)
        if name == "[Content_Types].xml":
            data = re.sub(rb'<Default Extension="jpeg" ContentType="image/jpeg"/>', b"", data)
            data = re.sub(rb'<Override PartName="/docProps/thumbnail.jpeg"[^>]*/>', b"", data)
        if name.startswith("ppt/embeddings/") and name.endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(data)) as e:
                inner = [(i.filename, e.read(i.filename)) for i in e.infolist()]
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w") as e:
                for n, d in inner:
                    if n == "docProps/core.xml":
                        d = re.sub(rb"<dc:creator>[^<]*</dc:creator>", b"<dc:creator>Laura Cohen</dc:creator>", d)
                        d = re.sub(rb"<cp:lastModifiedBy>[^<]*</cp:lastModifiedBy>",
                                   b"<cp:lastModifiedBy>Laura Cohen</cp:lastModifiedBy>", d)
                        d = re.sub(rb"(<dcterms:created[^>]*>)[^<]*", rb"\g<1>" + created.encode(), d)
                        d = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>" + modified.encode(), d)
                    if n == "docProps/app.xml":
                        d = re.sub(rb"<Company>[^<]*</Company>", b"<Company>Prairie Lantern Housing Fund</Company>", d)
                    zi = zipfile.ZipInfo(n, date_time=(1980, 1, 1, 0, 0, 0))
                    zi.compress_type = zipfile.ZIP_DEFLATED
                    e.writestr(zi, d)
            data = buf.getvalue()
        out.append((name, data))
    with zipfile.ZipFile(path, "w") as z:
        for name, data in out:
            zi = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, data)


def main(target, out_dir, figures_path):
    os.makedirs(out_dir, exist_ok=True)
    X = analyse(target)
    f, five, sixth = figures(X)
    write_xlsx(os.path.join(out_dir, XLSX), X, f, five, sixth)
    write_pptx(os.path.join(out_dir, PPTX), X, f, five, sixth)
    import postprocess as PP
    scrub_pptx(os.path.join(out_dir, PPTX), "2026-10-09T19:15:00Z", "2026-10-09T20:52:10Z")
    PP.normalize_ooxml(os.path.join(out_dir, XLSX), "Laura Cohen", "Laura Cohen", "2026-10-09T14:02:00Z",
                       "2026-10-09T20:41:37Z", excel=True)
    PP.normalize_ooxml(os.path.join(out_dir, PPTX), "Laura Cohen", "Laura Cohen", "2026-10-09T19:15:00Z",
                       "2026-10-09T20:52:10Z", title="Lenders to recruit first, 2027 rural program")
    stamp = dt.datetime(2026, 10, 9, 20, 52, 10, tzinfo=dt.timezone.utc).timestamp()
    for name in os.listdir(out_dir):
        os.utime(os.path.join(out_dir, name), (stamp, stamp))
    with open(figures_path, "w") as fh:
        json.dump(f, fh, indent=1, ensure_ascii=False)
        fh.write("\n")
    return X
