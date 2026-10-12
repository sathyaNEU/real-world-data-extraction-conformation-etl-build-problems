"""The pack's documents: the 2027 rural program guidelines (PDF), the recorder vendor's layout guide (PDF) and the
research desk's field dictionary (XLSX)."""
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                           # noqa: E402
from reportlab.lib.pagesizes import LETTER                                 # noqa: E402
from reportlab.lib.styles import ParagraphStyle                            # noqa: E402
from reportlab.lib.units import inch                                       # noqa: E402
from reportlab.pdfbase import pdfmetrics                                   # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                               # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import pack as P    # noqa: E402
import world as W   # noqa: E402

GUIDELINES = "rural_dpa_program_guidelines_2027.pdf"
LAYOUT = "venridge_pmm_layout_guide_v4.pdf"
DICTIONARY = "field_dictionary.xlsx"

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
pdfmetrics.registerFont(TTFont("LiberationSerif", FONT_DIR + "LiberationSerif-Regular.ttf"))
pdfmetrics.registerFont(TTFont("LiberationSerif-Bold", FONT_DIR + "LiberationSerif-Bold.ttf"))
pdfmetrics.registerFont(TTFont("LiberationSerif-Italic", FONT_DIR + "LiberationSerif-Italic.ttf"))
pdfmetrics.registerFont(TTFont("LiberationSans", FONT_DIR + "LiberationSans-Regular.ttf"))
pdfmetrics.registerFont(TTFont("LiberationSans-Bold", FONT_DIR + "LiberationSans-Bold.ttf"))
pdfmetrics.registerFont(TTFont("LiberationSans-Italic", FONT_DIR + "LiberationSans-Italic.ttf"))
INK = colors.HexColor("#1d1d1b")
GREY = colors.HexColor("#5f5f5a")
GREEN = colors.HexColor("#3b5b3a")
SLATE = colors.HexColor("#27455c")

# the fund's documents set in a serif; the vendor's in a sans
G_BODY = ParagraphStyle("gbody", fontName="LiberationSerif", fontSize=10.4, leading=14.2, textColor=INK, spaceAfter=5)
G_NUM = ParagraphStyle("gnum", parent=G_BODY, leftIndent=26, firstLineIndent=-26)
G_H1 = ParagraphStyle("gh1", fontName="LiberationSerif-Bold", fontSize=15.5, leading=19, textColor=GREEN, spaceAfter=3)
G_SUB = ParagraphStyle("gsub", fontName="LiberationSerif-Italic", fontSize=10, leading=13, textColor=GREY,
                       spaceAfter=10)
G_H2 = ParagraphStyle("gh2", fontName="LiberationSerif-Bold", fontSize=11.4, leading=15, textColor=GREEN, spaceBefore=8,
                      spaceAfter=3)
V_BODY = ParagraphStyle("vbody", fontName="LiberationSans", fontSize=9.2, leading=12.4, textColor=INK, spaceAfter=4.5)
V_H1 = ParagraphStyle("vh1", fontName="LiberationSans-Bold", fontSize=14, leading=17, textColor=SLATE, spaceAfter=2)
V_SUB = ParagraphStyle("vsub", fontName="LiberationSans", fontSize=9.4, leading=12, textColor=GREY, spaceAfter=9)
V_H2 = ParagraphStyle("vh2", fontName="LiberationSans-Bold", fontSize=10.4, leading=13.5, textColor=SLATE,
                      spaceBefore=8, spaceAfter=3)
CELL = ParagraphStyle("cell", fontName="LiberationSans", fontSize=8.2, leading=10.2, textColor=INK)
CELLB = ParagraphStyle("cellb", parent=CELL, fontName="LiberationSans-Bold")


def _table(rows, widths):
    data = [[Paragraph(str(c), CELLB if i == 0 else CELL) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, hAlign="LEFT", repeatRows=1)
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEABOVE", (0, 0), (-1, 0), 0.8, INK),
                           ("LINEBELOW", (0, 0), (-1, 0), 0.5, INK), ("LINEBELOW", (0, -1), (-1, -1), 0.8, INK),
                           ("TOPPADDING", (0, 0), (-1, -1), 1.8), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.8)]))
    return t


def _gframe(canvas, doc):
    canvas.saveState()
    canvas.setFont("LiberationSerif", 8)
    canvas.setFillColor(GREY)
    canvas.drawString(0.95 * inch, 0.6 * inch, "%s  |  Rural Down-Payment Assistance Program  |  Guidelines 2027" % W.ORG)
    canvas.drawRightString(LETTER[0] - 0.95 * inch, 0.6 * inch, "Page %d" % doc.page)
    canvas.restoreState()


def write_guidelines(path):
    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=0.95 * inch, rightMargin=0.95 * inch,
                            topMargin=0.9 * inch, bottomMargin=0.95 * inch,
                            title="Rural Down-Payment Assistance Program: Guidelines 2027", author=W.ORG, subject="",
                            creator="", producer="")
    s = [Paragraph("Rural Down-Payment Assistance Program", G_H1),
         Paragraph("Program guidelines for 2027. Approved by the board of directors on 18 June 2026; effective "
                   "1 January 2027. These guidelines govern the program in the rural service area only; the Wichita "
                   "and Topeka programs keep their own guidelines.", G_SUB)]
    N = lambda n, t: s.append(Paragraph("<b>%s</b>&nbsp;&nbsp;%s" % (n, t), G_NUM))
    s.append(Paragraph("1. Purpose", G_H2))
    N("1.1", "The fund helps households of modest income buy a first home. From 2027 it offers down-payment "
             "assistance in eight counties of central Kansas through banks, credit unions and mortgage companies that "
             "agree to originate program loans.")
    s.append(Paragraph("2. Service area", G_H2))
    N("2.1", "Saline, McPherson, Dickinson, Marion, Rice, Ottawa, Ellsworth and Lincoln counties.")
    N("2.2", "A home is in the county where it stands.")
    s.append(Paragraph("3. Buyers", G_H2))
    N("3.1", "A buyer qualifies who has not owned a home in the past three years, whose household income is at or "
             "below 100 per cent of the non-metropolitan Kansas median family income for the household's size, who "
             "will live in the home, and who has completed the fund's homebuyer education class.")
    N("3.2", "The buyer contributes at least $1,000 of their own funds at closing.")
    s.append(Paragraph("4. Assistance", G_H2))
    N("4.1", "Assistance is up to 4 per cent of the purchase price and no more than $12,500, as a deferred, "
             "zero-interest second mortgage forgiven in equal parts over five years of occupancy.")
    N("4.2", "The first mortgage is a fixed-rate, fully amortizing loan with a term of 15 to 30 years. Conventional, "
             "FHA, VA and USDA Rural Development loans qualify.")
    N("4.3", "Assistance is reserved through the fund's lender portal before closing and funded at closing.")
    s.append(Paragraph("5. Participating lenders", G_H2))
    N("5.1", "A participating lender signs the fund's participation agreement, takes the fund's lender training, "
             "and reserves, closes and reports program loans as the agreement provides.")
    N("5.2", "The board recruits participating lenders in rounds from a shortlist it agrees. Staff contact the "
             "lenders of each round in the order the board sets.")
    N("5.3", "A lender may decline or withdraw at any time; the next lender on the list is then approached.")
    s.append(Paragraph("6. Program year and reporting", G_H2))
    N("6.1", "The program year runs from 1 July to 30 June.")
    N("6.2", "Staff report reservations, assistance funded and home-purchase lending in the service area to the board "
             "by program year and by county.")
    s.append(Paragraph("7. Definitions", G_H2))
    N("7.1", "<i>Home</i>: a dwelling of one to four units, site-built or manufactured, on land in the service area.")
    N("7.2", "<i>Mortgage lender</i>: a bank, savings institution, credit union or licensed mortgage company. A seller "
             "who takes back a note from the buyer is not a mortgage lender.")
    N("7.3", "<i>Home-purchase mortgage</i>: a mortgage a mortgage lender makes to finance the purchase of a home in an "
             "arm's-length sale, secured by a first lien on the home, whatever the buyer's plans for living in it.")
    N("7.4", "<i>Arm's-length sale</i>: a sale between parties who are not related, on market terms. A transfer "
             "between a person and that person's trust, relative or business is not an arm's-length sale.")
    N("7.5", "A mortgage is <i>made</i> on the day it closes.")
    s.append(Spacer(1, 8))
    s.append(Paragraph("Questions on these guidelines go to the program director.", G_SUB))
    doc.build(s, onFirstPage=_gframe, onLaterPages=_gframe)


def _vframe(canvas, doc):
    canvas.saveState()
    canvas.setFont("LiberationSans", 7.4)
    canvas.setFillColor(GREY)
    canvas.drawString(0.8 * inch, 0.55 * inch, "Venridge Property Data, LLC  |  PMM file layout guide  |  v4, August 2026")
    canvas.drawRightString(LETTER[0] - 0.8 * inch, 0.55 * inch, "%d" % doc.page)
    canvas.restoreState()


V3_DESC = [
    ("REC_KEY", "Venridge record key", "text"),
    ("STATE", "State postal code", "text"),
    ("CNTY_FIPS", "County FIPS code of the recording county, the county in which the property lies", "5 digits"),
    ("CNTY_NAME", "Recording county (where the property lies)", "text"),
    ("MTG_DOC_NO", "Recorder's document number of the mortgage", "county format"),
    ("MTG_REC_DATE", "Date the mortgage was recorded", "MM/DD/YYYY"),
    ("MTG_INSTR_DATE", "Date on the mortgage instrument", "MM/DD/YYYY"),
    ("MTG_AMT", "Principal amount secured", "dollars"),
    ("MORTGAGEE_NAME", "Mortgagee as indexed by the recorder, first 80 characters", "text"),
    ("MORTGAGOR_NAME", "Mortgagor(s) as indexed", "text"),
    ("DEED_DOC_NO", "Document number of the matched conveyance", "county format"),
    ("DEED_REC_DATE", "Date the conveyance was recorded", "MM/DD/YYYY"),
    ("DEED_TYPE", "Conveyance type: WD warranty, JTWD joint tenancy warranty, SWD special warranty, TRD trustee's, "
                  "EXD executor's or administrator's, QCD quitclaim", "code"),
    ("DEED_CONSIDERATION", "Consideration as stated on the deed", "dollars or blank"),
    ("GRANTOR_NAME", "Grantor(s) of the conveyance", "text"),
    ("GRANTEE_NAME", "Grantee(s) of the conveyance", "text"),
    ("APN", "Parcel number (Kansas uniform parcel number)", "text"),
    ("SITUS_ADDR", "Property address, where the parcel has one", "text"),
    ("SITUS_CITY", "Postal city of the property address", "text"),
    ("SITUS_ZIP", "ZIP code of the property address", "5 digits"),
    ("LEGAL_BRIEF", "Abbreviated legal description", "text"),
    ("DOC_REMARKS", "Recorder's remarks on the mortgage", "text"),
    ("LOAD_DATE", "Date Venridge loaded the record", "MM/DD/YYYY"),
]
V4_CHANGES = [("MTG_REC_DATE", "RECORDING_DT"), ("MTG_INSTR_DATE", "DOC_DT"), ("MTG_AMT", "LOAN_AMT (whole dollars)"),
              ("MORTGAGEE_NAME", "LENDER_NAME"), ("MORTGAGOR_NAME", "BORROWER_NAME"),
              ("DEED_REC_DATE", "DEED_RECORDING_DT"), ("DEED_CONSIDERATION", "SALE_AMT"), ("GRANTOR_NAME", "SELLER_NAME"),
              ("GRANTEE_NAME", "BUYER_NAME"), ("DOC_REMARKS", "REMARKS"), ("LOAD_DATE", "LOAD_TS (date and time)"),
              ("STATE", "dropped"), ("(new)", "LOAN_TYPE_HINT: CONV, FHA, VA or USDA-RD where a rider names it")]


def write_layout(path):
    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                            topMargin=0.8 * inch, bottomMargin=0.85 * inch, title="PMM file layout guide v4",
                            author="Venridge Property Data, LLC", subject="", creator="", producer="")
    s = [Paragraph("Purchase-Money Mortgage (PMM) file: layout guide", V_H1),
         Paragraph("Version 4, August 2026. Applies to deliveries from September 2026; version 3 applies to earlier "
                   "deliveries.", V_SUB)]
    s.append(Paragraph("1. Contents", V_H2))
    s.append(Paragraph("A PMM file holds mortgages recorded in the subscribed counties that Venridge has matched to a "
                       "conveyance of the same property to the mortgagor. Each row is one recorded mortgage document "
                       "with the matched conveyance. Files are cut by the mortgage's recording date: an annual file "
                       "after each calendar year closes, and year-to-date files on request.", V_BODY))
    s.append(Paragraph("A mortgage that is re-recorded, for example to correct a legal description, appears again "
                       "under its new document number; the recorder's note is carried in DOC_REMARKS (REMARKS in "
                       "version 4).", V_BODY))
    s.append(Paragraph("2. Names and amounts", V_H2))
    s.append(Paragraph("Names are carried as the recorder indexed them, including nominee language. Kansas does not "
                       "require a deed to state the price paid, so DEED_CONSIDERATION (SALE_AMT) is the consideration "
                       "written on the deed and is often blank or nominal.", V_BODY))
    s.append(Paragraph("3. Version 3 layout", V_H2))
    s.append(Paragraph("Pipe-delimited text, one header row, Windows line endings.", V_BODY))
    s.append(_table([["Field", "Contents", "Format"]] + [list(r) for r in V3_DESC],
                    [1.45 * inch, 4.15 * inch, 1.3 * inch]))
    s.append(Paragraph("4. Changes in version 4", V_H2))
    s.append(Paragraph("Fields are reordered with the property fields first; dates are written YYYY-MM-DD. Renamed "
                       "and new fields:", V_BODY))
    s.append(_table([["Version 3", "Version 4"]] + [list(r) for r in V4_CHANGES], [1.9 * inch, 4.9 * inch]))
    s.append(Paragraph("5. Support", V_H2))
    s.append(Paragraph("Data questions to the Venridge data desk; delivery questions to your account manager.", V_BODY))
    doc.build(s, onFirstPage=_vframe, onLaterPages=_vframe)


HMDA_FIELDS = [
    ("activity_year", "Calendar year of the action taken"), ("lei", "Legal entity identifier of the reporting institution"),
    ("derived_msa-md", "MSA or metropolitan division of the property; 99999 outside any"),
    ("state_code", "State of the property"), ("county_code", "Five-digit FIPS code of the county the reporting institution's geocoding assigned to the property address"),
    ("census_tract", "Eleven-digit census tract of the property"),
    ("conforming_loan_limit", "C conforming, NC nonconforming, U undetermined, NA"),
    ("derived_loan_product_type", "Loan type and lien status"),
    ("derived_dwelling_category", "Construction method and number of units"),
    ("derived_ethnicity", "Derived ethnicity of applicants"), ("derived_race", "Derived race of applicants"),
    ("derived_sex", "Derived sex of applicants"),
    ("action_taken", "1 loan originated; 2 approved, not accepted; 3 denied; 4 withdrawn; 5 closed for incompleteness; "
                     "6 purchased loan; 7 preapproval request denied; 8 preapproval request approved, not accepted"),
    ("purchaser_type", "Type of entity that bought the loan in the year; 0 not sold"),
    ("preapproval", "1 requested, 2 not requested"), ("loan_type", "1 conventional, 2 FHA, 3 VA, 4 USDA RHS or FSA"),
    ("loan_purpose", "1 home purchase, 2 home improvement, 31 refinancing, 32 cash-out refinancing, 4 other, 5 NA"),
    ("lien_status", "1 first lien, 2 subordinate lien"),
    ("reverse_mortgage", "1 reverse mortgage, 2 not, 1111 exempt"),
    ("open-end_line_of_credit", "1 open-end, 2 not, 1111 exempt"),
    ("business_or_commercial_purpose", "1 primarily business or commercial, 2 not, 1111 exempt"),
    ("loan_amount", "Amount, at the midpoint of its $10,000 interval"),
    ("combined_loan_to_value_ratio", "Combined loan-to-value ratio, per cent"),
    ("interest_rate", "Interest rate, per cent"), ("rate_spread", "Spread over the average prime offer rate"),
    ("hoepa_status", "1 high-cost, 2 not high-cost, 3 not applicable"),
    ("total_loan_costs", "Total loan costs, dollars"), ("origination_charges", "Total origination charges, dollars"),
    ("loan_term", "Months"), ("property_value", "Value relied on, at the midpoint of its $10,000 interval"),
    ("construction_method", "1 site-built, 2 manufactured home"),
    ("occupancy_type", "1 principal residence, 2 second residence, 3 investment property"),
    ("total_units", "Dwelling units"), ("income", "Gross annual income relied on, thousands of dollars"),
    ("debt_to_income_ratio", "Debt-to-income ratio, banded"), ("applicant_age", "Age of the applicant, banded"),
    ("submission_of_application", "1 submitted directly to the institution, 2 not, 3 not applicable, 1111 exempt"),
    ("initially_payable_to_institution", "1 initially payable to the institution, 2 not, 3 not applicable, 1111 exempt"),
    ("aus-1", "First automated underwriting system used; 6 not applicable, 1111 exempt"),
    ("denial_reason-1", "First reason for denial; 10 not applicable, 1111 exempt"),
    ("tract_population", "Census tract population"),
    ("tract_minority_population_percent", "Minority share of the tract population, per cent"),
    ("ffiec_msa_md_median_family_income", "FFIEC median family income of the MSA or non-MSA area"),
    ("tract_to_msa_income_percentage", "Tract median family income against the area's, per cent"),
    ("tract_owner_occupied_units", "Owner-occupied units in the tract"),
    ("tract_one_to_four_family_homes", "One-to-four family homes in the tract"),
    ("tract_median_age_of_housing_units", "Median age of the tract's housing units, years")]
PANEL_FIELDS = [("activity_year", "Reporting year"), ("lei", "Legal entity identifier"),
                ("agency_code", "1 OCC, 2 FRS, 3 FDIC, 5 NCUA, 7 HUD, 9 CFPB"), ("id_rssd", "Federal Reserve RSSD ID; -1 none"),
                ("respondent_name", "Institution name as filed"), ("respondent_city", "City of the institution"),
                ("respondent_state", "State of the institution"), ("parent_name", "Parent company, where one is filed"),
                ("assets", "Total assets, thousands of dollars (depositories)"),
                ("other_lender_code", "0 depository or independent; other codes for subsidiaries and affiliates")]
OWN_FIELDS = [
    (P.FILES["deliveries"], "delivery_id", "Venridge delivery identifier"),
    (P.FILES["deliveries"], "file_name", "File as received"), (P.FILES["deliveries"], "layout", "PMM layout version"),
    (P.FILES["deliveries"], "recordings_from", "First recording date covered"),
    (P.FILES["deliveries"], "recordings_to", "Last recording date covered"),
    (P.FILES["deliveries"], "rows", "Data rows in the file"), (P.FILES["deliveries"], "delivered", "Date received"),
    (P.FILES["alias"], "mortgagee_as_indexed", "Mortgagee text exactly as it appears in a Venridge file"),
    (P.FILES["alias"], "nmls_id", "NMLS ID of the lender the text refers to"),
    (P.FILES["alias"], "lender_name", "Legal name of that lender"),
    (P.FILES["alias"], "match", "exact, variant, or MERS as nominee for the lender"),
    (P.FILES["alias"], "first_seen", "First recording date the text was seen"),
    (P.FILES["alias"], "reviewed", "Date the match was reviewed"),
    (P.FILES["register"], "nmls_id", "NMLS ID"), (P.FILES["register"], "legal_name", "Legal name"),
    (P.FILES["register"], "institution_type", "Bank, credit union or mortgage company"),
    (P.FILES["register"], "charter", "Charter or licence"), (P.FILES["register"], "lei", "Legal entity identifier, where one is held"),
    (P.FILES["register"], "rssd_id", "Federal Reserve RSSD ID, where one is held"),
    (P.FILES["register"], "service_area_offices", "Offices in the eight counties"),
    (P.FILES["agg2023"], "originations", "Loans originated in 2023"),
    (P.FILES["agg2023"], "loan_amount_000", "Their total amount, thousands of dollars"),
    (P.FILES["education"], "completions", "Households completing the class"),
    (P.FILES["education"], "first_time_buyers", "Of which, households who have not owned a home in three years")]
GRAIN = [
    (P.FILES["lar"], "One row per application, preapproval request or purchased loan an institution reported for a "
                     "Kansas property; the public file carries no loan number, address or action date."),
    (P.FILES["panel"], "One row per reporting institution and year."),
    ("venridge_pmm_ks8_*.txt", "One row per recorded mortgage document Venridge matched to a conveyance; files cut by "
                               "the mortgage's recording date."),
    (P.FILES["deliveries"], "One row per Venridge delivery."),
    (P.FILES["alias"], "One row per mortgagee text seen in a Venridge file and matched to a lender."),
    (P.FILES["register"], "One row per lender known to the fund."),
    (P.FILES["shortlist"], "One row per lender on the board's 2027 shortlist."),
    (P.FILES["agg2023"], "One row per county, loan purpose and loan type with any 2023 origination."),
    (P.FILES["education"], "One row per month, county and class format with any completion."),
]


def write_dictionary(path):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fields"
    ws.append(["File", "Field", "Meaning"])
    for f, d in HMDA_FIELDS:
        ws.append([P.FILES["lar"], f, d])
    for f, d in PANEL_FIELDS:
        ws.append([P.FILES["panel"], f, d])
    for f, ds, _ in V3_DESC:
        ws.append(["venridge_pmm_ks8_2024.txt, _2025.txt", f, ds])
    ws.append(["venridge_pmm_ks8_2026ytd.txt", "(all)", "Layout v4; see the layout guide for renamed fields"])
    for f, fld, d in OWN_FIELDS:
        ws.append([f, fld, d])
    ws2 = wb.create_sheet("Grain")
    ws2.append(["File", "Grain"])
    for f, g in GRAIN:
        ws2.append([f, g])
    for sh, widths in ((ws, (40, 34, 96)), (ws2, (40, 110))):
        for c in sh[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="3B5B3A")
        for col, wd in zip("ABC", widths):
            sh.column_dimensions[col].width = wd
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        sh.freeze_panes = "A2"
    wb.save(path)
