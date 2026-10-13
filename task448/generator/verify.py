"""Independent verifier for task448: reads only the shipped bundle with DuckDB, on its own code path (no calc.py, no
generator state), recomputes each shortlisted lender's count, the five, the sixth and every county figure, and
compares them with the golden workbook, the golden deck and figures.json.

    python3 task448/generator/verify.py [target] [golden]
"""
import json
import math
import os
import sys

import duckdb

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
FIPS = ("20169", "20113", "20041", "20115", "20159", "20143", "20053", "20105")
FOUR = ("SALINE", "MCPHERSON", "DICKINSON", "MARION")


def rnd(x):
    import math
    return int(math.floor(x / 10.0 + 0.5) * 10)


def main(target=None, golden=None):
    target = target or os.path.join(TASK, "target")
    golden = golden or os.path.join(TASK, "golden")
    q = duckdb.connect()
    t = lambda f: os.path.join(target, f).replace("'", "''")
    q.execute("""CREATE TABLE hm AS SELECT lei, count(*) AS n FROM read_csv('%s', all_varchar=true)
                 WHERE county_code IN %s AND action_taken = '1' AND loan_purpose = '1' AND lien_status = '1'
                 GROUP BY lei""" % (t("hmda_lar_kansas_2024_2025.csv"), str(FIPS)))
    v3 = """SELECT CNTY_NAME AS county, MTG_DOC_NO AS mdoc, strptime(MTG_REC_DATE, '%%m/%%d/%%Y')::DATE AS mrec,
              strptime(MTG_INSTR_DATE, '%%m/%%d/%%Y')::DATE AS instr, CAST(MTG_AMT AS DOUBLE) AS amt,
              trim(MORTGAGEE_NAME) AS mtgee, MORTGAGOR_NAME AS mtgor, DEED_DOC_NO AS ddoc,
              strptime(DEED_REC_DATE, '%%m/%%d/%%Y')::DATE AS drec, DEED_TYPE AS dtype, GRANTOR_NAME AS grantor,
              coalesce(SITUS_ADDR, '') AS situs, coalesce(DOC_REMARKS, '') AS remarks, '%s' AS src
              FROM read_csv('%s', delim='|', header=true, all_varchar=true, quote='')"""
    v4 = """SELECT CNTY_NAME, MTG_DOC_NO, CAST(RECORDING_DT AS DATE), CAST(DOC_DT AS DATE), CAST(LOAN_AMT AS DOUBLE),
              trim(LENDER_NAME), BORROWER_NAME, DEED_DOC_NO, CAST(DEED_RECORDING_DT AS DATE), DEED_TYPE, SELLER_NAME,
              coalesce(SITUS_ADDR, ''), coalesce(REMARKS, ''), '%s'
              FROM read_csv('%s', delim='|', header=true, all_varchar=true, quote='')"""
    q.execute("CREATE TABLE f AS " + v3 % ("2024", t("venridge_pmm_ks8_2024.txt")) + " UNION ALL " +
              v3 % ("2025", t("venridge_pmm_ks8_2025.txt")) + " UNION ALL " +
              v4 % ("2026", t("venridge_pmm_ks8_2026ytd.txt")))
    q.execute("""CREATE TABLE a AS SELECT trim(mortgagee_as_indexed) AS mtgee, CAST(nmls_id AS VARCHAR) AS nmls
                 FROM read_csv('%s', all_varchar=true)""" % t("mortgagee_alias_register.csv"))
    q.execute("""CREATE TABLE r AS SELECT CAST(nmls_id AS VARCHAR) AS nmls, legal_name, coalesce(lei, '') AS lei
                 FROM read_csv('%s', all_varchar=true)""" % t("lender_register_2026-09.csv"))
    q.execute("""CREATE TABLE fr AS SELECT f.*, a.nmls, f.mrec = f.drec AS packaged,
                 f.dtype = 'QCD' AS qcd, (f.remarks LIKE '%RE-RECORD%' OR f.remarks LIKE '%RERECORD%') AS rerec
                 FROM f LEFT JOIN a USING (mtgee)""")
    # the larger mortgage against a conveyance is the first lien
    q.execute("""CREATE TABLE f1 AS SELECT * FROM (SELECT *, row_number() OVER (PARTITION BY county, ddoc
                 ORDER BY amt DESC, mdoc) AS k FROM fr WHERE NOT qcd AND NOT rerec) WHERE k = 1""")
    # one farmstead mortgage recorded in two counties: keep the county with the house
    q.execute("""CREATE TABLE mtg AS SELECT * FROM f1 WHERE nmls IS NOT NULL AND NOT (
                 situs = '' AND (nmls, mtgor, amt, instr) IN (SELECT (nmls, mtgor, amt, instr) FROM f1
                 GROUP BY ALL HAVING count(DISTINCT county) > 1))""")
    q.execute("""CREATE TABLE lend AS SELECT nmls, count(*) AS rows, avg(packaged::INT) AS pk FROM mtg
                 WHERE instr BETWEEN DATE '2024-01-01' AND DATE '2025-12-31' GROUP BY nmls""")
    q.execute("""CREATE TABLE rep AS SELECT r.nmls, hm.n, coalesce(l.rows, 0) AS rows,
                 CASE WHEN coalesce(l.pk, 1) >= 0.5 THEN 'P' ELSE 'S' END AS route
                 FROM r JOIN hm ON r.lei = hm.lei LEFT JOIN lend l ON l.nmls = r.nmls""")
    cap = dict(q.execute("SELECT route, sum(rows) / sum(n) FROM rep GROUP BY route").fetchall())
    import openpyxl
    sl = openpyxl.load_workbook(os.path.join(target, "rural_lender_shortlist_2027.xlsx"), read_only=True)["Shortlist"]
    short = [(str(r[4]), r[1]) for r in sl.iter_rows(min_row=2, values_only=True)]
    hmn = dict(q.execute("SELECT nmls, n FROM rep").fetchall())
    lend = {n: (rows, pk) for n, rows, pk in q.execute("SELECT nmls, rows, pk FROM lend").fetchall()}
    counts = {}
    for n, name in short:
        if n in hmn:
            counts[name] = float(hmn[n])
        else:
            rows, pk = lend.get(n, (0, 1.0))
            counts[name] = rows / cap["P" if pk >= 0.5 else "S"]
    order = sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))
    # SF11: the vendor files that the file is not every purchase-money mortgage, and its rows alone rank wrongly
    from pypdf import PdfReader
    guide = " ".join(" ".join(pg.extract_text() for pg in PdfReader(
        os.path.join(target, "venridge_pmm_layout_guide_v4.pdf")).pages).split()).lower()
    rows_only = sorted(((lend.get(n, (0, 1.0))[0], name) for n, name in short), reverse=True)
    sf11 = ("does not hold every purchase-money mortgage recorded in a county" in guide and
            {nm for _, nm in rows_only[:5]} != {nm for nm, _ in order[:5]})
    # county figures: reporters' HMDA count spread by their recorded counties; others over their route's capture
    two = {}
    for c, n, k, route, is_rep, tot in q.execute("""
            WITH m AS (SELECT * FROM mtg WHERE instr BETWEEN DATE '2024-01-01' AND DATE '2025-12-31'),
            bc AS (SELECT nmls, county, count(*) AS k FROM m GROUP BY ALL),
            lr AS (SELECT nmls, sum(k) AS t FROM bc GROUP BY nmls)
            SELECT bc.county, bc.nmls, bc.k, CASE WHEN coalesce(l.pk, 1) >= 0.5 THEN 'P' ELSE 'S' END,
                   rep.nmls IS NOT NULL, lr.t
            FROM bc JOIN lr USING (nmls) LEFT JOIN rep USING (nmls) LEFT JOIN lend l USING (nmls)""").fetchall():
        add = hmn[n] * k / tot if is_rep else k / cap[route]
        key = c if c in FOUR else "OTHER"
        two[key] = two.get(key, 0.0) + add
    seen = {x[0] for x in q.execute("SELECT DISTINCT nmls FROM mtg WHERE instr BETWEEN DATE '2024-01-01' AND "
                                    "DATE '2025-12-31'").fetchall()}
    for n in hmn:
        if n not in seen:
            raise SystemExit("VERIFY: a reporter with no recorded mortgage needs HMDA's county codes")
    # SF12: guidelines 2.2 place a home in the county in which the property is recorded; HMDA's county is a geocode.
    # Keeping the recorded mortgages where recorded and only the rest of each reporter's count where HMDA's codes put
    # it lands in the same bins; HMDA's codes for every reporter loan do not.
    name_of = dict(zip(FIPS, ("SALINE", "MCPHERSON", "DICKINSON", "MARION", "RICE", "OTTAWA", "ELLSWORTH", "LINCOLN")))
    rules = " ".join(" ".join(pg.extract_text() for pg in PdfReader(
        os.path.join(target, "rural_dpa_program_guidelines_2027.pdf")).pages).split()).lower()
    sf12_pin = "a home is in the county where it stands, which is the county in which the property is recorded" in rules
    geo = {}
    for lei, cc, k in q.execute("""SELECT lei, county_code, count(*) FROM read_csv('%s', all_varchar=true)
            WHERE county_code IN %s AND action_taken = '1' AND loan_purpose = '1' AND lien_status = '1'
            GROUP BY ALL""" % (t("hmda_lar_kansas_2024_2025.csv"), str(FIPS))).fetchall():
        geo.setdefault(lei, {})[name_of[cc]] = k
    lei_of = dict(q.execute("SELECT nmls, lei FROM r").fetchall())
    hy, hc = {}, {}
    for c, n, k, route, is_rep, tot in q.execute("""
            WITH m AS (SELECT * FROM mtg WHERE instr BETWEEN DATE '2024-01-01' AND DATE '2025-12-31'),
            bc AS (SELECT nmls, county, count(*) AS k FROM m GROUP BY ALL),
            lr AS (SELECT nmls, sum(k) AS t FROM bc GROUP BY nmls)
            SELECT bc.county, bc.nmls, bc.k, CASE WHEN coalesce(l.pk, 1) >= 0.5 THEN 'P' ELSE 'S' END,
                   rep.nmls IS NOT NULL, lr.t
            FROM bc JOIN lr USING (nmls) LEFT JOIN rep USING (nmls) LEFT JOIN lend l USING (nmls)""").fetchall():
        key = c if c in FOUR else "OTHER"
        if is_rep:
            hy[key] = hy.get(key, 0.0) + k
        else:
            hy[key] = hy.get(key, 0.0) + k / cap[route]
            hc[key] = hc.get(key, 0.0) + k / cap[route]
    for n in hmn:
        g = geo.get(lei_of[n], {})
        rec = sum(x[0] for x in q.execute("""SELECT count(*) FROM mtg WHERE nmls = ? AND instr BETWEEN
                  DATE '2024-01-01' AND DATE '2025-12-31'""", [n]).fetchall())
        for c, k in g.items():
            key = c if c in FOUR else "OTHER"
            hy[key] = hy.get(key, 0.0) + (hmn[n] - rec) * k / hmn[n]
            hc[key] = hc.get(key, 0.0) + k
    sf12 = (sf12_pin and all(rnd(hy[k]) == rnd(two[k]) for k in two) and
            sum(rnd(hc[k]) != rnd(two[k]) for k in two) >= 4)
    print("verify: SF12 hybrid %s; HMDA's codes %s" % (", ".join("%s %.1f" % (k, hy[k]) for k in two),
                                                       ", ".join("%s %.1f" % (k, hc[k]) for k in two)))
    py = {}
    for c, packaged, k in q.execute("""SELECT county, packaged, count(*) FROM mtg
            WHERE instr BETWEEN DATE '2025-07-01' AND DATE '2026-06-30' GROUP BY ALL""").fetchall():
        key = c if c in FOUR else "OTHER"
        py[key] = py.get(key, 0.0) + k / cap["P" if packaged else "S"]
    py["TOTAL"] = sum(py.values())
    # SF13: guidelines 7.5 count every mortgage a lender makes, including any no report or data file shows
    sf13 = "every mortgage a lender makes is counted, including any that no report or data file shows" in rules
    # SF14: the completed lenders at a capture measured within each county and route land in the same bins
    lr_tot = dict(q.execute("""SELECT nmls, count(*) FROM mtg WHERE instr BETWEEN DATE '2024-01-01' AND
                  DATE '2025-12-31' GROUP BY nmls""").fetchall())
    byc = q.execute("""SELECT nmls, county, count(*) FROM mtg WHERE instr BETWEEN DATE '2024-01-01' AND
                  DATE '2025-12-31' GROUP BY ALL""").fetchall()
    rt = {n: ("P" if pk >= 0.5 else "S") for n, (_, pk) in lend.items()}
    cn, cd = {}, {}
    for n, c, k in byc:
        if n in hmn:
            cn[(rt[n], c)] = cn.get((rt[n], c), 0.0) + k
            cd[(rt[n], c)] = cd.get((rt[n], c), 0.0) + hmn[n] * k / lr_tot[n]
    cr = {}
    for n, c, k in byc:
        key = c if c in FOUR else "OTHER"
        cr[key] = cr.get(key, 0.0) + (hmn[n] * k / lr_tot[n] if n in hmn else k / (cn[(rt[n], c)] / cd[(rt[n], c)])
                                      if cd.get((rt[n], c)) else k / cap[rt[n]])
    sf14 = all(rnd(cr[k]) == rnd(two[k]) for k in two)
    # SF15: the reporters' July to December 2025 from their 2025 HMDA counts lands in the same bins
    h25 = dict(q.execute("""SELECT r.nmls, count(*) FROM read_csv('%s', all_varchar=true) l JOIN r ON r.lei = l.lei
              WHERE county_code IN %s AND action_taken = '1' AND loan_purpose = '1' AND lien_status = '1'
              AND activity_year = '2025' GROUP BY ALL""" % (t("hmda_lar_kansas_2024_2025.csv"), str(FIPS))).fetchall())
    y25 = dict(q.execute("""SELECT nmls, count(*) FROM mtg WHERE instr BETWEEN DATE '2025-01-01' AND
              DATE '2025-12-31' GROUP BY nmls""").fetchall())
    hh = {}
    for n, c, packaged, instr in q.execute("""SELECT nmls, county, packaged, instr FROM mtg
            WHERE instr BETWEEN DATE '2025-07-01' AND DATE '2026-06-30'""").fetchall():
        key = c if c in FOUR else "OTHER"
        if n in hmn and str(instr) <= "2025-12-31":
            hh[key] = hh.get(key, 0.0) + h25.get(n, 0) / y25[n]
        else:
            hh[key] = hh.get(key, 0.0) + 1.0 / cap["P" if packaged else "S"]
    hh["TOTAL"] = sum(hh.values())
    rk = lambda k, x: (int(math.floor(x / 100.0 + 0.5) * 100) if k == "TOTAL" else rnd(x))
    sf15 = all(rk(k, hh[k]) == rk(k, py[k]) for k in py)
    # SF16: the unweighted mean of the reporters' own ratios within each group lands Ask A in the same bins
    ratios = {}
    for n, hmc, rows_, route_ in q.execute("SELECT nmls, n, rows, route FROM rep").fetchall():
        ratios.setdefault(route_, []).append(rows_ / hmc)
    capm = {k: sum(v) / len(v) for k, v in ratios.items()}
    pm = {}
    for c, packaged, k in q.execute("""SELECT county, packaged, count(*) FROM mtg
            WHERE instr BETWEEN DATE '2025-07-01' AND DATE '2026-06-30' GROUP BY ALL""").fetchall():
        key = c if c in FOUR else "OTHER"
        pm[key] = pm.get(key, 0.0) + k / capm["P" if packaged else "S"]
    pm["TOTAL"] = sum(pm.values())
    sf16 = all(rk(k, pm[k]) == rk(k, py[k]) for k in py)
    print("verify: SF14 county-and-route capture %s; SF15 HMDA half-year %s" % (
        ", ".join("%s %.1f" % (k, cr[k]) for k in two), ", ".join("%s %.1f" % (k, hh[k]) for k in py)))
    # the write-up's path: the rows the stack leaves out, recomputed
    st = dict(zip(("stacked", "window", "rerec", "qcd"), q.execute("""SELECT count(*),
              count(*) FILTER (WHERE instr BETWEEN DATE '2024-01-01' AND DATE '2025-12-31'),
              count(*) FILTER (WHERE rerec AND instr BETWEEN DATE '2024-01-01' AND DATE '2025-12-31'),
              count(*) FILTER (WHERE qcd AND NOT rerec AND instr BETWEEN DATE '2024-01-01' AND DATE '2025-12-31')
              FROM fr""").fetchone()))
    win = "instr BETWEEN DATE '2024-01-01' AND DATE '2025-12-31'"
    st["second"] = q.execute("SELECT count(*) FROM fr WHERE NOT qcd AND NOT rerec AND %s" % win).fetchone()[0] - \
        q.execute("SELECT count(*) FROM f1 WHERE %s" % win).fetchone()[0]
    st["mortgage"] = q.execute("SELECT count(*) FROM mtg WHERE %s" % win).fetchone()[0]
    st["seller"] = q.execute("SELECT count(*) FROM f1 WHERE nmls IS NULL AND %s" % win).fetchone()[0]
    st["land"] = q.execute("SELECT count(*) FROM f1 WHERE nmls IS NOT NULL AND %s" % win).fetchone()[0] - st["mortgage"]
    # compare with the goldens
    wb = openpyxl.load_workbook(os.path.join(golden, "rural_lender_recruitment_2027.xlsx"), data_only=True)
    ws = wb["Recruitment order"]
    rows = [(ws.cell(r, 2).value, ws.cell(r, 3).value) for r in range(5, 17)]
    fail = []
    if [n for n, _ in rows] != [n for n, _ in order]:
        fail.append("order")
    for name, v in rows:
        if v != rnd(counts[name]):
            fail.append("count %s" % name)
    wc = wb["County volumes"]
    for i, k in enumerate(FOUR + ("OTHER",)):
        if wc.cell(4 + i, 2).value != rnd(two[k]) or wc.cell(4 + i, 3).value != rnd(py[k]):
            fail.append("county %s" % k)
    if wc.cell(9, 3).value != rk("TOTAL", py["TOTAL"]):
        fail.append("program-year total")
    from pptx import Presentation
    ch = [s for s in Presentation(os.path.join(golden, "rural_lender_recruitment_2027.pptx")).slides[0].shapes
          if s.has_chart][0].chart
    chart = dict(zip(ch.plots[0].categories, ch.plots[0].series[0].values))
    if sorted(chart.values()) != sorted(rnd(v) for v in counts.values()):
        fail.append("chart")
    if not sf11:
        fail.append("SF11: the guide's coverage sentence, or the file's rows alone giving a wrong five")
    if not sf12:
        fail.append("SF12: guidelines 2.2's recorded county, the hybrid's bins, or HMDA's codes missing four figures")
    fig = json.load(open(os.path.join(HERE, "figures.json")))
    fs = fig["stack"]
    want = {"stacked": fs["stacked"], "window": fs["closed_2024_2025"], "rerec": fs["rerecordings"],
            "qcd": fs["quitclaim_pairings"], "second": fs["second_mortgages"], "land": fs["land_only_records"],
            "seller": fs["seller_notes"], "mortgage": fs["lender_mortgages"]}
    bad = [k for k, v in want.items() if int(v.replace(",", "")) != st[k]]
    if bad:
        fail.append("the write-up's stack %s" % ", ".join("%s %s against %s" % (k, st[k], want[k]) for k in bad))
    if not sf13:
        fail.append("SF13: guidelines 7.5's count of every mortgage a lender makes")
    if not sf14:
        fail.append("SF14: the county-and-route capture leaves the county figures' bins")
    if not sf15:
        fail.append("SF15: the reporters' HMDA half-year leaves the program year's bins")
    if not sf16:
        fail.append("SF16: the mean of the reporters' ratios leaves the program year's bins")
    if fig["capture_packaged"] != "%d" % round(100 * cap["P"]) or fig["capture_separate"] != "%d" % round(100 * cap["S"]):
        fail.append("captures in figures.json")
    print("verify: capture %.4f packaged, %.4f separate; five %s; sixth %s %.1f" % (
        cap["P"], cap["S"], " | ".join("%s %.1f" % (n.split(",")[0], v) for n, v in order[:5]), order[5][0],
        order[5][1]))
    print("verify: 2024-25 %s; program year %s" % (
        ", ".join("%s %.1f" % kv for kv in two.items()), ", ".join("%s %.1f" % kv for kv in py.items())))
    if fail:
        raise SystemExit("VERIFY FAILED: " + "; ".join(fail))
    print("verify: PASS (bundle-only recomputation agrees with the workbook, the deck and figures.json)")


if __name__ == "__main__":
    main(*(sys.argv[1:3]))
