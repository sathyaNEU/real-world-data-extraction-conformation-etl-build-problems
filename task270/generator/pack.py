"""Assemble the task270 evidence pack: the world, the records in their systems' layouts, the documents."""
import os

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font

import common as K
import docs as D
import records as RC
import world as W


def register_xlsx(path, packs, eq):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Packs"
    ws.append(["Pack", "Store", "Store name", "Refrigerant", "Full charge (kg)", "In service from", "Decommissioned",
               "Leak detection"])
    for r in packs:
        ws.append([None if v == "" else v for v in r])
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, wd in zip("ABCDEFGH", (11, 7, 26, 12, 15, 15, 15, 14)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A2"
    w2 = wb.create_sheet("Cabinets and cold rooms")
    w2.append(["Asset", "Store", "Type", "Refrigerant", "Charge (kg)", "Installed"])
    for r in eq:
        w2.append(r)
    for c in w2[1]:
        c.font = Font(bold=True)
    for col, wd in zip("ABCDEF", (9, 7, 17, 12, 12, 10)):
        w2.column_dimensions[col].width = wd
    w2.freeze_panes = "A2"
    w3 = wb.create_sheet("Notes")
    w3["A1"] = "F-gas register kept by the property team; charges as commissioned or as last certified by Kelbray."
    w3["A1"].alignment = Alignment(wrap_text=True)
    w3.column_dimensions["A"].width = 90
    wb.save(path)


def kpi_stats(w):
    jb = w["jobs"]
    out = {}
    rng = np.random.default_rng([K.SEED, 59])
    for reg in K.REGIONS:
        sub = jb[(jb.pack.map(w["packs"].region) == reg) & (jb.date.map(lambda d: d.year) == K.YEAR)]
        out[reg] = dict(routine=int((sub.type == "ROUTINE").sum()), breakdown=int((sub.type == "BREAKDOWN").sum()),
                        hours=int(rng.integers(5, 14)), fix=int(rng.integers(78, 93)))
    return out


def build_pack(out, seed=K.SEED):
    os.makedirs(out, exist_ok=True)
    w = W.extras(W.simulate(seed), seed)
    rc = RC.Records(w, seed)
    p = lambda k: os.path.join(out, K.FILES[k])   # noqa: E731
    rc.stores(p("stores"))
    packs, eq = rc.register_rows()
    register_xlsx(p("register"), packs, eq)
    rc.counts["register"] = len(packs) + len(eq)
    rc.jobs(p("jobs"), p("cylinders"))
    rc.levels(p("levels"), p("tags"))
    rc.power(p("power"))
    D.strategy(p("strategy"))
    D.spec(p("spec"))
    D.bureau(p("bureau"))
    D.kpi(p("kpi"), kpi_stats(w))
    D.dictionary(p("dictionary"))
    D.contents(p("contents"), rc.counts)
    return w, rc.counts
