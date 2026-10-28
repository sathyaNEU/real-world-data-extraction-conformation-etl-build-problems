"""task270 workbooks: the ABS estimated resident population by area, sex and five-year group, the other states'
notifications of residents' deaths, and the region's residential aged care services."""
from __future__ import annotations

import datetime as dt

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font

import common as K

GROUPS = [(lo, lo + 4) for lo in range(0, 85, 5)] + [(85, 115)]
GROUP_LABELS = ["%d-%d" % g for g in GROUPS[:-1]] + ["85 and over"]
BOLD = Font(bold=True)


def five_year(v):
    return [int(v[lo:hi + 1].sum()) for lo, hi in GROUPS]


def published(groups, y, area, sex):
    """The Bureau's estimate of each group: the population with the estimate's own error, about 0.4 per cent of a
    group (the components it is built from are themselves estimated)."""
    out = []
    for j, n in enumerate(groups):
        e = K.rng("erp_error", y, area, sex, j).normal(0, 0.004)
        out.append(int(round(n * (1 + e))))
    return out


def erp_workbook(E, A, path):
    """Regional population by age and sex: one sheet per 30 June, males, females and persons for each area."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contents"
    lines = ["Australian Bureau of Statistics",
             "Regional population by age and sex, 2026",
             "Estimated resident population by age and sex, local government areas (2023 boundaries), "
             "selected areas of New South Wales, 30 June 2016 to 30 June 2026",
             "Released 27 August 2027",
             "",
             "Tables",
             ]
    for j, t in enumerate(lines, 1):
        ws.cell(row=j, column=1, value=t)
        if j in (1, 2):
            ws.cell(row=j, column=1).font = BOLD
    for k, y in enumerate(range(2016, 2027)):
        ws.cell(row=7 + k, column=1, value=str(y))
        ws.cell(row=7 + k, column=2, value="Estimated resident population, 30 June %d%s" % (
            y, " (preliminary)" if y == 2026 else (" (revised)" if y >= 2022 else "")))
    ws.cell(row=19, column=1, value="Estimates for 2022 to 2025 have been revised and estimates for 2026 are "
                                    "preliminary.")
    ws.cell(row=21, column=1, value="(c) Commonwealth of Australia 2027")
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 70
    for y in range(2016, 2027):
        sh = wb.create_sheet(str(y))
        sh.cell(row=1, column=1, value="Estimated resident population by age and sex, 30 June %d%s" % (
            y, " (preliminary)" if y == 2026 else ""))
        sh.cell(row=1, column=1).font = BOLD
        head = ["LGA code", "Local government area", "Sex"] + GROUP_LABELS + ["Total"]
        for c, h in enumerate(head, 1):
            sh.cell(row=3, column=c, value=h).font = BOLD
        r = 4
        for a in sorted(A, key=lambda a: a.code):
            m = published(five_year(E[(y, a.name, "M")]), y, a.name, "M")
            f = published(five_year(E[(y, a.name, "F")]), y, a.name, "F")
            for sex, vals in (("Males", m), ("Females", f), ("Persons", [x + z for x, z in zip(m, f)])):
                row = [a.code, a.name, sex] + vals + [sum(vals)]
                for c, v in enumerate(row, 1):
                    sh.cell(row=r, column=c, value=v)
                r += 1
        sh.column_dimensions["B"].width = 22
        for c in range(4, 4 + len(head)):
            sh.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 10
    wb.save(path)


def interstate_workbook(rows, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Deaths"
    ws.cell(row=1, column=1, value="Deaths of usual residents of the Tidewell Coast and Ranges PHN region registered "
                                   "in other states and territories, registered 1 January 2016 to 30 June 2027")
    ws.cell(row=1, column=1).font = BOLD
    ws.cell(row=2, column=1, value="Supplied by the Australian Bureau of Statistics from the national deaths "
                                   "collection, 6 September 2027")
    head = ["State of registration", "Registration number", "Date of registration", "Date of death", "Date of birth",
            "Sex", "Age at death", "Usual residence SA2 code (ASGS 2021)", "Usual residence SA2 name"]
    for c, h in enumerate(head, 1):
        ws.cell(row=4, column=c, value=h).font = BOLD
    for r, row in enumerate(rows, 5):
        for c, v in enumerate(row, 1):
            if c in (3, 4, 5):
                v = dt.date.fromisoformat(v)
            cell = ws.cell(row=r, column=c, value=v)
            if c in (3, 4, 5):
                cell.number_format = "dd/mm/yyyy"
    for c, w in zip("ABCDEFGHI", (12, 18, 14, 14, 14, 7, 8, 18, 26)):
        ws.column_dimensions[c].width = w
    wb.save(path)


RACF_SUFFIX = ["Aged Care", "Lodge", "Gardens", "Residential Care", "Court", "House", "Manor", "Village"]
PROVIDERS = ["Tidewell Aged Care Ltd", "Coast and Ranges Care Co-operative Ltd", "Ranges Community Living Ltd",
             "Bannockvale District Benevolent Society", "Shoreline Seniors Care Pty Ltd", "Coastal Care Alliance Ltd",
             "Inland Rivers Aged Care Ltd", "Saltwick Shoals Community Hostel Inc", "Hessian Flat Pioneers Home Inc",
             "Kestrelmoor Ageing Well Ltd", "Bluewater Lifestyle Care Pty Ltd", "Ranges Health Services Ltd"]


def racf_workbook(A, P_old_85, path):
    """Residential aged care services in the region at 30 June 2026 (the department's service list, filtered by the
    network to its SA2s)."""
    g = K.rng("racf")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Services"
    ws.cell(row=1, column=1, value="Residential aged care services, Tidewell Coast and Ranges PHN region, "
                                   "30 June 2026").font = BOLD
    head = ["Service name", "Approved provider", "Physical address", "SA2 code (ASGS 2021)", "SA2 name",
            "Care type", "Residential places", "Restorative care places", "Organisation type"]
    for c, h in enumerate(head, 1):
        ws.cell(row=3, column=c, value=h).font = BOLD
    r = 4
    streets = ["Hospital Road", "Church Street", "Ocean Parade", "Railway Street", "Showground Road", "Banksia Drive",
               "Wattle Street", "River Road", "Station Street", "Park Avenue", "Mill Lane", "High Street"]
    used = set()
    streets = streets + ["Beach Road", "Bridge Street", "Gibson Avenue", "Kurrawa Street", "Lighthouse Road",
                         "Marine Parade", "Old Coast Road", "Pacific Street", "Pioneer Avenue", "Quarry Road",
                         "Reservoir Road", "School Lane", "Seaview Terrace", "Shiers Avenue", "Telegraph Road",
                         "Tidewell Highway", "Union Street", "Victoria Street", "Waratah Drive", "Wharf Street"]
    for a in A:
        need = P_old_85[a.name] * 0.19
        k = max(1, int(round(need / 72)))
        sizes = np.maximum(24, np.round(g.dirichlet(np.ones(k) * 2.5) * need)).astype(int)
        w = np.array([3.0 if "Surrounds" not in nm else 0.6 for _, nm in a.sa2])
        for j in range(k):
            sa2 = a.sa2[int(g.choice(len(a.sa2), p=w / w.sum()))]
            town = sa2[1].replace(" Surrounds", "").split(" - ")[0]
            for t in range(40):
                suffix = RACF_SUFFIX[int(g.integers(0, len(RACF_SUFFIX)))]
                name = "%s %s" % (town, suffix) if t < 20 else "%s %s %s" % (
                    town, ["North", "South", "East", "West", "Park", "Grove"][t % 6], suffix)
                if name not in used:
                    break
            used.add(name)
            prov = PROVIDERS[int(g.integers(0, len(PROVIDERS)))]
            orgt = "Not-for-profit" if "Pty" not in prov else "Private"
            row = [name, prov, "%d %s" % (int(g.integers(2, 240)), streets[int(g.integers(0, len(streets)))]),
                   sa2[0], sa2[1], "Residential", int(sizes[j]), int(g.choice([0, 0, 0, 1, 2, 3, 4])), orgt]
            for c, v in enumerate(row, 1):
                ws.cell(row=r, column=c, value=v)
            r += 1
    for c, w in zip("ABCDEFGHI", (30, 40, 24, 18, 26, 12, 12, 12, 16)):
        ws.column_dimensions[c].width = w
    wb.save(path)
