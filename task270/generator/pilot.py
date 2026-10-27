"""The 2025-26 pilot in Krathwold and Lathcombe: every consultation the pilot pharmacies held from 1 October 2025 to
30 September 2026 and the follow-ups after it, the services portal's export as pulled on 2 October 2026, the federation's
monthly payment statements, and the community pharmacy list.

The portal labels each record with the network on the pharmacy's portal profile as it stands when the export is
pulled. A pharmacy claims a consultation through the portal by the fifth working day of the month after it, and the
next month's statement pays it, so the September consultations recorded after the pull, and the follow-ups after it,
reach only the statements. The pharmacies' patient record systems post each consultation to the portal, and a
pharmacist sometimes also keys the same first consultation in by hand, typing its date a day or three out and on a day
of their own; both are paid, and the federation recovers the hand-keyed claim's payment in the statement after the one
that paid it. A pharmacy that claims after the fifth working day is paid all the same, and the federation recovers that
payment in the statement after too: the consultation was held, the claim was late."""
from __future__ import annotations

import datetime as dt
import os

import openpyxl
from openpyxl.styles import Font

from common import BANK_HOLIDAYS, NETWORK_NAME, REKEY, TARGET, rng, token, write_csv

PILOT_START = dt.date(2025, 10, 1)
PILOT_END = dt.date(2026, 9, 30)
PORTAL_PULL = dt.date(2026, 10, 2)
RELOCATE = dt.date(2026, 6, 1)
PILOT_NETS = ("KRT", "LTH")
FEE = {"First consultation": 28.0, "Follow-up": 14.0}

F_PORTAL = "services_portal_statin_pilot_export_pulled_2026-10-02.csv"
F_STATEMENTS = "pilot_payment_statements_2025-11_to_2027-01.xlsx"
F_PHARMACIES = "community_pharmacies_2026-12.csv"


def quarter(d):
    """The pilot's quarters: 1 October to 31 December 2025 is Q1, then Q2 to Q4 to 30 September 2026."""
    m = (d.year - 2025) * 12 + d.month - 10
    return 1 + m // 3


def _workday(d):
    while d.weekday() == 6:
        d += dt.timedelta(days=1)
    return d


def deadline(d):
    """The claim's last day: the fifth working day of the month after the consultation."""
    t = dt.date(d.year + (d.month == 12), d.month % 12 + 1, 1)
    n = 0
    while True:
        if t.weekday() < 5 and t not in BANK_HOLIDAYS:
            n += 1
            if n == 5:
                return t
        t += dt.timedelta(days=1)


def _off_six_weeks(d3, first):
    """No follow-up falls exactly six weeks after either record of its first consultation, so a reading of "within six
    weeks" that stops at 41 days and one that stops at 42 count the same patients."""
    marks = {first["date"], first.get("dup_date") or first["date"]}
    while any((d3 - m).days == 42 for m in marks) or d3.weekday() == 6:
        d3 += dt.timedelta(days=1)
    return d3


def _on_time(d, delay):
    return min(d + dt.timedelta(days=delay), deadline(d))


def build(pharms, pts):
    r = rng("pilot")
    pilots = [p for p in pharms if p["kind"] == "community" and p["network"] in PILOT_NETS]
    moved = [p for p in pilots if p["network"] == "KRT"][1]
    moved["relocated_on"] = RELOCATE
    moved["previous_network"] = "KRT"
    moved["network_now"] = "MRL"
    by_ph = {}
    for pt in pts:
        if pt["kind"] != "sm" or pt["left"] or pt["band"] != "65+":
            continue
        if pt["network"] not in PILOT_NETS:
            continue
        usual = max(sorted(set(x[1] for x in pt["items"])), key=lambda o: sum(1 for x in pt["items"] if x[1] == o)) \
            if pt["items"] else None
        by_ph.setdefault(usual, []).append(pt)
    events = []           # (true record) pharmacy, patient, type, consultation date, recorded on, duplicate of
    n_days = (PILOT_END - PILOT_START).days + 1
    for ph in pilots:
        cand = by_ph.get(ph["ods"], [])
        k = int(round(len(cand) * float(r.uniform(0.16, 0.24))))
        chosen = list(r.choice(len(cand), size=min(k, len(cand)), replace=False)) if cand else []
        for i in chosen:
            pt = cand[i]
            d1 = _workday(PILOT_START + dt.timedelta(days=int(r.integers(0, n_days))))
            if d1 > PILOT_END:
                d1 = PILOT_END - dt.timedelta(days=1)
            first = dict(ph=ph["ods"], pt=pt, type="First consultation", date=d1,
                         rec=_on_time(d1, int(r.choice([0, 0, 0, 1, 1, 2, 3, 4, 6, 9]))), dup_of=None)
            events.append(first)
            if r.random() < 0.05:            # keyed in again by hand, with the date retyped
                _dup(r, events, first, ph, pt)
            if r.random() < 0.74:
                k = (first["dup_date"] - d1).days if first.get("dup_date") else 0
                if k >= 2 and r.random() < 0.6:
                    # the follow-up falls just past six weeks from the consultation, inside six weeks of the copy
                    gap = 42 + int(r.integers(1, k))
                else:
                    gap = int(r.choice([14, 18, 21, 24, 28, 31, 35, 38, 41, 44, 45, 49, 56]))
                d3 = _off_six_weeks(_workday(d1 + dt.timedelta(days=gap)), first)
                events.append(dict(ph=ph["ods"], pt=pt, type="Follow-up", date=d3,
                                   rec=_on_time(d3, int(r.choice([0, 0, 1, 1, 2, 5]))), dup_of=None))
    net_of = {p["ods"]: p["network"] for p in pilots}

    def net_on(ods, d):
        return "MRL" if ods == moved["ods"] and d >= RELOCATE else net_of[ods]

    _aim_boundaries(r, events, net_on)
    _late_claims(r, events, net_on)
    for e in events:
        e["id"] = "SP" + token("pilot", e["pt"]["pid"], e["type"], e["date"].isoformat(), e["rec"].isoformat(),
                               n=8).upper()
        e["key"] = e["pt"]["old_key"] if e["date"] < REKEY else e["pt"]["new_key"]
    return pilots, moved, events


def _dup(r, events, first, ph, pt, shift=None):
    """A first consultation keyed in by hand as well: its date typed a day or three out either way, keyed on a day of
    the pharmacist's own, before or after the system's record reached the portal."""
    d1 = first["date"]
    if shift is None:
        shift = int(r.choice([-3, -2, -1, 1, 2, 3]))
    d2 = d1 + dt.timedelta(days=shift)
    if d2.weekday() == 6:
        d2 += dt.timedelta(days=1 if shift > 0 else -1)
    rec = min(max(d1, d2) + dt.timedelta(days=int(r.integers(0, 8))), deadline(d2))
    if rec == first["rec"]:
        rec += dt.timedelta(days=1 if rec < deadline(d2) else -1)
    events.append(dict(ph=ph if isinstance(ph, str) else ph["ods"], pt=pt, type="First consultation", date=d2,
                       rec=max(rec, d2), dup_of=first))
    first["dup_date"] = d2


BOUNDS = (dt.date(2026, 1, 1), dt.date(2026, 4, 1), dt.date(2026, 7, 1))
EARLIER = (1, 2, 1)     # retyped copies dated before a boundary for a consultation just after it, per boundary
LATER = (2, 1, 2)       # retyped copies dated after a boundary for a consultation just before it


def _aim_boundaries(r, events, net_on):
    """In each pilot network, first consultations within a few days of each internal quarter boundary get a
    retyped copy on the other side of it, in counts that differ boundary to boundary. Each such patient's follow-up
    falls three to four weeks after the consultation."""
    firsts = [e for e in events if e["type"] == "First consultation" and e["dup_of"] is None and not e.get("dup_date")]
    for net in PILOT_NETS:
        for b, ne, nl in zip(BOUNDS, EARLIER, LATER):
            after = sorted((e for e in firsts if net_on(e["ph"], e["date"]) == net and 0 <= (e["date"] - b).days <= 4
                            and not e.get("dup_date")), key=lambda e: ((e["date"] - b).days, e["pt"]["pid"]))
            before = sorted((e for e in firsts if net_on(e["ph"], e["date"]) == net and 1 <= (b - e["date"]).days <= 4
                             and not e.get("dup_date")), key=lambda e: ((b - e["date"]).days, e["pt"]["pid"]))
            for e in after[:ne]:
                _dup(r, events, e, e["ph"], e["pt"], shift=-((e["date"] - b).days + int(r.integers(1, 3))))
                _follow(r, events, e)
            for e in before[:nl]:
                _dup(r, events, e, e["ph"], e["pt"], shift=(b - e["date"]).days + int(r.integers(0, 2)))
                _follow(r, events, e)


def _follow(r, events, first):
    """The aimed consultation's follow-up, three to four weeks on, replacing any drawn for it."""
    pid = first["pt"]["pid"]
    events[:] = [x for x in events if not (x["type"] == "Follow-up" and x["pt"]["pid"] == pid)]
    d3 = _off_six_weeks(_workday(first["date"] + dt.timedelta(days=int(r.integers(21, 29)))), first)
    events.append(dict(ph=first["ph"], pt=first["pt"], type="Follow-up", date=d3, rec=d3, dup_of=None))


LAST_LATE = dt.date(2026, 11, 30)     # a late claim recorded by then is paid in December and recovered in January
LATE_FIRSTS = {"KRT": (2, 1, 3, 2), "LTH": (1, 3, 2, 2)}   # late-claimed first consultations, per pilot quarter
LATE_FOLLOW = {"KRT": (1, 0, 1, 2), "LTH": (2, 1, 0, 1)}   # late-claimed follow-ups, by the first consultation's quarter


def _late_claims(r, events, net_on):
    """Some consultations are claimed after the fifth working day of the month after them: paid, then recovered.
    Each pilot network and quarter gets late-claimed first consultations whose patient was followed up within six
    weeks, and late-claimed follow-ups, in counts that differ quarter to quarter; none is a recorded-twice patient."""
    twice = {e["pt"]["pid"] for e in events if e["dup_of"] is not None}
    fu = {}
    for e in events:
        if e["type"] == "Follow-up":
            fu.setdefault(e["pt"]["pid"], []).append(e)
    firsts = [e for e in events if e["type"] == "First consultation" and e["dup_of"] is None
              and e["pt"]["pid"] not in twice and PILOT_START <= e["date"] <= PILOT_END]
    firsts.sort(key=lambda e: (e["date"], e["pt"]["pid"]))
    for net in PILOT_NETS:
        for q in (1, 2, 3, 4):
            pool = [e for e in firsts if net_on(e["ph"], e["date"]) == net and quarter(e["date"]) == q
                    and any(0 < (f["date"] - e["date"]).days <= 42 for f in fu.get(e["pt"]["pid"], []))]
            idx = list(r.choice(len(pool), size=LATE_FIRSTS[net][q - 1] + LATE_FOLLOW[net][q - 1], replace=False))
            for k, i in enumerate(idx):
                e = pool[i]
                if k < LATE_FIRSTS[net][q - 1]:
                    e["rec"] = min(deadline(e["date"]) + dt.timedelta(days=int(r.integers(3, 36))), LAST_LATE)
                    e["late"] = True
                else:
                    f = min((f for f in fu[e["pt"]["pid"]] if 0 < (f["date"] - e["date"]).days <= 42),
                            key=lambda f: f["date"])
                    f["rec"] = min(deadline(f["date"]) + dt.timedelta(days=int(r.integers(3, 36))), LAST_LATE)
                    f["late"] = True


def truth(events, pilots, moved):
    """The golden: each patient's first consultation once, at its own date, in the network the pharmacy was in that
    day; each followed up within six weeks of that date by any pilot pharmacy."""
    net_of = {p["ods"]: p["network"] for p in pilots}

    def net_on(ods, d):
        if ods == moved["ods"] and d >= RELOCATE:
            return "MRL"
        return net_of[ods]

    firsts = {}
    for e in events:
        if e["type"] == "First consultation" and e["dup_of"] is None and PILOT_START <= e["date"] <= PILOT_END:
            firsts[e["pt"]["pid"]] = e
    fups = {}
    for e in events:
        if e["type"] == "Follow-up":
            fups.setdefault(e["pt"]["pid"], []).append(e["date"])
    C1 = {(n, q): 0 for n in PILOT_NETS for q in (1, 2, 3, 4)}
    C2 = dict(C1)
    for pid, e in firsts.items():
        n = net_on(e["ph"], e["date"])
        if n not in PILOT_NETS:
            continue
        q = quarter(e["date"])
        C1[(n, q)] += 1
        if any(0 < (d - e["date"]).days <= 42 for d in fups.get(pid, [])):
            C2[(n, q)] += 1
    return C1, C2


def write_all(pilots, moved, events, pharms):
    os.makedirs(TARGET, exist_ok=True)
    name = {p["ods"]: p["name"] for p in pharms}
    portal_net = {p["ods"]: NETWORK_NAME[p.get("network_now", p["network"])] for p in pilots}
    rows = []
    for e in sorted(events, key=lambda x: (x["rec"], x["ph"], x["id"])):
        if e["rec"] > PORTAL_PULL:
            continue
        rows.append(dict(record_id=e["id"], pharmacy_ods=e["ph"], pharmacy_name=name[e["ph"]],
                         network=portal_net[e["ph"]], patient_ref=e["key"], consultation_type=e["type"],
                         consultation_date=e["date"].strftime("%d/%m/%Y"), recorded_on=e["rec"].strftime("%d/%m/%Y"),
                         claim_status="Paid" if e["rec"] < dt.date(2026, 9, 1) else "Claimed"))
    n_portal = write_csv(os.path.join(TARGET, F_PORTAL), list(rows[0].keys()), rows)
    # statements: a month's statement pays what was recorded in the month before; a duplicate first consultation's
    # payment is recovered in the statement after the one that paid it
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    months = []
    m = dt.date(2025, 11, 1)
    while m <= dt.date(2027, 1, 1):
        months.append(m)
        m = dt.date(m.year + (m.month == 12), m.month % 12 + 1, 1)
    lines = {mm: [] for mm in months}
    for e in events:
        paid = dt.date(e["rec"].year + (e["rec"].month == 12), e["rec"].month % 12 + 1, 1)
        if paid in lines:
            lines[paid].append((e["ph"], e["id"], e["key"], e["type"], e["date"], FEE[e["type"]], "Payment", ""))
        if e["dup_of"] is not None or e.get("late"):
            rec_m = dt.date(paid.year + (paid.month == 12), paid.month % 12 + 1, 1)
            if rec_m in lines:
                lines[rec_m].append((e["ph"], e["id"], e["key"], e["type"], e["date"], -FEE[e["type"]], "Recovery",
                                     ""))
    for mm in months:
        ws = wb.create_sheet(mm.strftime("%b %Y"))
        ws.append(["Pharmacy ODS", "Record", "Patient", "Service", "Date of service", "Amount (£)", "Line"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for row in sorted(lines[mm], key=lambda x: (x[0], x[4], x[1])):
            ws.append([row[0], row[1], row[2], row[3], row[4], row[5], row[6]])
        for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
            for c in row:
                c.number_format = "dd/mm/yyyy"
        for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
            for c in row:
                c.number_format = "#,##0.00"
        for col, w in zip("ABCDEFG", (13, 12, 16, 20, 15, 12, 10)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
    wb.save(os.path.join(TARGET, F_STATEMENTS))
    # the community pharmacy list
    out = []
    for p in pharms:
        out.append(dict(ods_code=p["ods"], pharmacy_name=p["name"], owner=p["owner"] or "Independent",
                        contract_type="Distance selling" if p["kind"] == "distance-selling" else "Community pharmacy",
                        locality=NETWORK_NAME[p.get("network_now", p["network"])] if p["network"] else "",
                        relocated_on=p["relocated_on"].isoformat() if p.get("relocated_on") else "",
                        previous_locality=NETWORK_NAME[p["previous_network"]] if p.get("previous_network") else ""))
    out.sort(key=lambda x: x["ods_code"])
    n_ph = write_csv(os.path.join(TARGET, F_PHARMACIES), list(out[0].keys()), out)
    return {F_PORTAL: n_portal, F_STATEMENTS: sum(len(v) for v in lines.values()), F_PHARMACIES: n_ph}
