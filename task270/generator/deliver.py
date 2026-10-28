"""task270 goldens: the planning team's area workbook and the executive paper's chart, both computed from the shipped
files by calc.py, and figures.json, from which the write-up is drafted."""
from __future__ import annotations

import json
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt                       # noqa: E402
from matplotlib import font_manager                    # noqa: E402
import openpyxl                                        # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402

import calc                                            # noqa: E402
import common as K                                     # noqa: E402
import postprocess as PP                               # noqa: E402

INK = "#1f2a30"
MUTED = "#5d6870"
QUAL = "#2a78d6"       # the program's areas
REST = "#b9c2c8"       # the other areas
LINE_C = "#1f2a30"
PREPARED = "2 October 2027"


def figures(target):
    out, deaths, C, E = calc.rungs(target)
    by_code, _ = calc.read_corr(target)
    rows = calc.read_ah(target)
    visits, residents = calc.ah_golden(rows, by_code)
    areas = sorted(K.AREA_NAMES)
    e85 = {a: round(out[a]["G"], 1) for a in areas}
    raw = {a: out[a]["G"] for a in areas}
    call = sorted(a for a in areas if raw[a] < K.LINE)
    above = sorted((a for a in areas if raw[a] >= K.LINE), key=lambda a: raw[a])
    d85 = {a: int(out[a]["D"].sum()) for a in areas}
    F = {
        "call": call,
        "call_count": len(call),
        "line": K.LINE,
        "e85": e85,
        "qualifies": {a: ("Yes" if a in call else "No") for a in areas},
        "nearest_above": above[0],
        "nearest_above_e85": e85[above[0]],
        "highest": max(areas, key=lambda a: raw[a]),
        "highest_e85": e85[max(areas, key=lambda a: raw[a])],
        "lowest": min(areas, key=lambda a: raw[a]),
        "lowest_e85": e85[min(areas, key=lambda a: raw[a])],
        "deaths_85plus_2024_2026": d85,
        "deaths_85plus_total": sum(d85.values()),
        "visits": visits,
        "residents": residents,
        "default_split_e85": {a: round(out[a]["R2"], 1) for a in areas},
        "default_split_call": sorted(a for a in areas if out[a]["R2"] < K.LINE),
        # the filed terms and conventions the write-up states, read from the build's own parameters
        "conventions": {"from_age": 85, "single_to": K.TOP - 1, "open_from": K.TOP, "late_days": 41,
                        "year_days": 365, "late_from_day_of_july": 1, "late_to_day_of_august": 10,
                        "clauses": ["3.1", "3.2", "3.3", "3.4"]},
    }
    return F, out


def workbook(F, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Area list"
    bold = Font(name="Calibri", size=11, bold=True, color="1F2A30")
    head_fill = PatternFill("solid", fgColor="E3EEF1")
    qual_fill = PatternFill("solid", fgColor="DCEBFA")
    thin = Side(style="thin", color="B9C6CC")
    ws["A1"] = "Late-life frailty program 2028 to 2030: area list"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F2A30")
    ws["A2"] = ("%d of the 12 areas qualify: %s. Life expectancy at 85 under each area's 2024 to 2026 death rates, "
                "persons (program guidelines 3.1 to 3.4); after-hours home visits to residents aged 85 and over in "
                "2026." % (F["call_count"], ", ".join(F["call"][:-1]) + " and " + F["call"][-1]))
    ws["A2"].font = Font(name="Calibri", size=10, color="5D6870")
    head = ["LGA code", "Local government area", "Life expectancy at 85 (years)", "Qualifies (below 7.0 years)",
            "After-hours visits to residents aged 85+, 2026", "Residents aged 85+ visited, 2026"]
    for c, h in enumerate(head, 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = bold
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="bottom", horizontal="left" if c <= 2 else "right")
        cell.border = Border(bottom=thin)
    ws.row_dimensions[4].height = 45
    r = 5
    for a in sorted(K.AREA_NAMES, key=lambda a: K.AREA_CODE[a]):
        vals = [K.AREA_CODE[a], a, F["e85"][a], F["qualifies"][a], F["visits"][a], F["residents"][a]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if c == 3:
                cell.number_format = "0.0"
            if c in (5, 6):
                cell.number_format = "#,##0"
            if c >= 3:
                cell.alignment = Alignment(horizontal="right")
            if F["qualifies"][a] == "Yes":
                cell.fill = qual_fill
        r += 1
    ws.cell(row=r, column=2, value="Twelve areas").font = bold
    for c, key in ((5, "visits"), (6, "residents")):
        cell = ws.cell(row=r, column=c, value=sum(F[key].values()))
        cell.number_format = "#,##0"
        cell.font = bold
        cell.alignment = Alignment(horizontal="right")
        cell.border = Border(top=thin)
    ws.cell(row=r + 2, column=1, value="Areas on their 1 July 2023 boundaries. Qualifying areas shaded. An area is "
                                       "named on its unrounded life expectancy (guidelines 3.4).").font = \
        Font(name="Calibri", size=9, color="5D6870")
    for col, w in zip("ABCDEF", (10, 24, 16, 16, 20, 18)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C5"
    ws.print_title_rows = "4:4"

    ns = wb.create_sheet("Notes")
    notes = [
        ("Life expectancy at 85", "Years a person reaching 85 would live on average if the area's 2024 to 2026 "
                                  "death rate at each age held: a life table by single year of age from 85 to 94 "
                                  "with 95 and over as the last group, closed on its own rate, half a year lived in "
                                  "the year of death (guidelines 3.2)."),
        ("Deaths", "Deaths of the area's usual residents at each age (completed years) that occurred in 2024, 2025 or "
                   "2026: the Registry's extract (registrations to 30 June 2027) and the ABS file of residents' "
                   "deaths registered in other states, each resident's area from the usual residence's 2021 SA2 by "
                   "the ABS correspondence to 2023 areas."),
        ("Population at each age", "The ABS publishes each area's estimated resident population for 85 and over as "
                                   "one total. Each age's population at 30 June 2024, 2025 and 2026 is rebuilt from "
                                   "the 2021 census counts by single year (SA2s of usual residence summed to areas): "
                                   "each census age followed forward as the people of its birth window, the 41 of "
                                   "365 days born 1 July to 10 August a year younger at 30 June than the rest, less "
                                   "the deaths of that window since census night, each year's 85-and-over total then "
                                   "brought to the published estimate."),
        ("Why not the census mix", "Splitting each area's 85-and-over estimate by its 2021 census ages puts the "
                                   "coast's cohorts that turned 85 after census night outside the group and leaves "
                                   "Spraymouth's late-eighties cohorts where they were in 2021: Saltwick Shoals and "
                                   "Dunesford read above seven years that way and Spraymouth below."),
        ("After-hours visits", "Times the after-hours service attended a resident aged 85 and over at the time, by the "
                               "arrival's date in 2026: clinicians who arrived together count once and a later return "
                               "on the same case counts again (the export has one row per clinician attending)."),
        ("Residents visited", "Different residents aged 85 and over visited in 2026: from 1 July 2026 the service's "
                              "new system gave carried-over clients new keys and kept the January to June key in its "
                              "external identifier, so each resident is counted once across the two halves."),
        ("Prepared", "Population Health Planning, %s. Files taken 15 September 2027." % PREPARED),
    ]
    ns["A1"] = "Notes"
    ns["A1"].font = Font(name="Calibri", size=13, bold=True, color="1F2A30")
    for j, (k, v) in enumerate(notes, 3):
        ns.cell(row=j, column=1, value=k).font = bold
        cell = ns.cell(row=j, column=2, value=v)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ns.cell(row=j, column=1).alignment = Alignment(vertical="top")
        ns.row_dimensions[j].height = 15 * (1 + len(v) // 95)
    ns.column_dimensions["A"].width = 24
    ns.column_dimensions["B"].width = 100
    wb.save(path)


def chart(F, path):
    fp = "/usr/share/fonts/opentype/inter/Inter-Regular.otf"
    fb = "/usr/share/fonts/opentype/inter/Inter-SemiBold.otf"
    for f in (fp, fb):
        if os.path.exists(f):
            font_manager.fontManager.addfont(f)
    plt.rcParams["font.family"] = "Inter" if os.path.exists(fp) else "DejaVu Sans"
    plt.rcParams["svg.hashsalt"] = "task270"
    areas = sorted(K.AREA_NAMES, key=lambda a: (F["e85"][a], a), reverse=True)
    vals = [F["e85"][a] for a in areas]
    colors = [QUAL if F["qualifies"][a] == "Yes" else REST for a in areas]
    fig, ax = plt.subplots(figsize=(9.6, 6.2), dpi=150)
    y = range(len(areas))
    ax.barh(list(y), vals, color=colors, height=0.62, edgecolor="white", linewidth=1.5)
    ax.axvline(K.LINE, color=LINE_C, linewidth=1.2, linestyle=(0, (4, 3)))
    ax.text(K.LINE - 0.06, -0.85, "Program line, 7.0 years", fontsize=9, color=INK, va="center", ha="right")
    for i, (a, v) in enumerate(zip(areas, vals)):
        if F["qualifies"][a] == "Yes":
            ax.text(v - 0.06, i, "%.1f" % v, va="center", ha="right", fontsize=9.5, color="white",
                    fontweight="semibold")
            ax.text(K.LINE + 0.08, i, "qualifies", va="center", fontsize=9.5, color=INK, fontweight="semibold")
        else:
            ax.text(v + 0.06, i, "%.1f" % v, va="center", fontsize=9.5, color=INK)
    ax.set_ylim(-1.3, len(areas) - 0.4)
    ax.set_yticks(list(y))
    ax.set_yticklabels(areas, fontsize=10, color=INK)
    ax.set_xlim(0, 8.6)
    ax.set_xticks([0, 1, 2, 3, 4, 5, 6, 7, 8])
    ax.set_xlabel("Life expectancy at 85, years (2024 to 2026 death rates, persons)", fontsize=9.5, color=MUTED)
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.spines["bottom"].set_color("#b9c6cc")
    ax.tick_params(axis="x", colors=MUTED, labelsize=9)
    ax.tick_params(axis="y", length=0)
    ax.grid(axis="x", color="#e6ebee", linewidth=0.6)
    ax.set_axisbelow(True)
    names = F["call"]
    title = "%s, %s, %s and %s fall below seven years" % tuple(
        sorted(names, key=lambda a: F["e85"][a])) if len(names) == 4 else "Areas below seven years"
    fig.suptitle(title, x=0.015, y=0.975, ha="left", fontsize=13.5, fontweight="semibold", color=INK)
    fig.text(0.015, 0.925, "Late-life frailty program 2028 to 2030: life expectancy at 85 by local government area",
             fontsize=10, color=MUTED, ha="left")
    fig.text(0.015, 0.018, "Deaths: NSW Registry of Births, Deaths and Marriages (DRA-2025-117) and ABS notifications "
                           "of deaths registered interstate. Populations: ABS estimated resident population, rebuilt "
                           "by single year from the 2021 Census. Tidewell Coast and Ranges PHN, %s." % PREPARED,
             fontsize=7.5, color=MUTED, ha="left", wrap=True)
    fig.subplots_adjust(left=0.17, right=0.97, top=0.88, bottom=0.14)
    fig.savefig(path, dpi=150, metadata={"Software": None, "Title": title, "Author": K.PHN})
    plt.close(fig)


def run(task_dir, target):
    gold = os.path.join(task_dir, "golden")
    os.makedirs(gold, exist_ok=True)
    F, out = figures(target)
    workbook(F, os.path.join(gold, K.G_BOOK))
    PP.normalize_ooxml(os.path.join(gold, K.G_BOOK), K.PEOPLE["planning"], K.PEOPLE["planning"],
                       "2027-09-29T23:10:00Z", "2027-10-01T22:45:00Z", excel=True, company=K.PHN)
    chart(F, os.path.join(gold, K.G_CHART))
    with open(os.path.join(task_dir, "generator", "figures.json"), "w") as fh:
        json.dump(F, fh, indent=1, sort_keys=True)
    return F
