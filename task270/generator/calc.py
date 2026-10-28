"""task270 calculations, read from the shipped files only: each area's life expectancy at 85 on the golden's
construction and on every rung of the ladder, the after-hours asks, and the cracker's and mirror's paths.

The golden rebuilds each area's population at each age from 85 to 94 (and 95 and over) at 30 June 2024, 2025 and
2026 from the 2021 census counts by single year of age, each census age followed forward as the people of its birth
window (a ninth of whom, born 1 July to 10 August, are a year younger at 30 June than the rest), less the deaths
the registry and the other states record for that window since census night, and brings each year's rebuilt total
to the area's estimated resident population aged 85 and over. The rates are the deaths at each age in 2024 to 2026
over those exposures, and the life table runs by single year to 94 with 95 and over closed on its own rate."""
from __future__ import annotations

import csv
import datetime as dt
import os
from collections import defaultdict

import numpy as np
import openpyxl

import common as K

AGES = list(range(85, K.TOP + 1))          # 85..94 single, 95 = 95 and over
NA = len(AGES)
CEN = K.CENSUS


def _d(s):
    return dt.date.fromisoformat(s)


def age_on(dob, day):
    return day.year - dob.year - ((day.month, day.day) < (dob.month, dob.day))


# ------------------------------------------------------------------------------------------------- readers
def read_corr(target):
    by_code, by_name = {}, {}
    with open(os.path.join(target, K.F_CORR), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            by_code[r["SA2_CODE_2021"]] = r["LGA_NAME_2023"]
            by_name[r["SA2_NAME_2021"]] = r["SA2_CODE_2021"]
    return by_code, by_name


def read_deaths(target, by_code):
    """Every death of a usual resident: the registry's extract and the other states' notifications. Each row:
    (area, sex, dob or None, dod, age at death)."""
    out = []
    with open(os.path.join(target, K.F_DEATHS), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            area = by_code.get(r["usual_residence_sa2_code"])
            if area is None:
                continue
            dob = _d(r["date_of_birth"]) if r["date_of_birth"] else None
            out.append((area, r["sex"], dob, _d(r["date_of_death"]), int(r["age_at_death"]), "NSW"))
    wb = openpyxl.load_workbook(os.path.join(target, K.F_INTERSTATE), read_only=True)
    ws = wb["Deaths"]
    head = None
    for row in ws.iter_rows(values_only=True):
        if head is None:
            if row and row[0] == "State of registration":
                head = list(row)
            continue
        if row[0] is None:
            continue
        rec = dict(zip(head, row))
        area = by_code.get(str(rec["Usual residence SA2 code (ASGS 2021)"]))
        if area is None:
            continue
        dob = rec["Date of birth"]
        dod = rec["Date of death"]
        dob = dob.date() if hasattr(dob, "date") else dob
        dod = dod.date() if hasattr(dod, "date") else dod
        out.append((area, rec["Sex"], dob, dod, int(rec["Age at death"]), rec["State of registration"]))
    wb.close()
    return out


def read_census(target, by_name, by_code):
    """Census counts by area (2023 boundaries), sex and single age 0..115 from the TableBuilder export."""
    out = defaultdict(lambda: np.zeros(116))
    with open(os.path.join(target, K.F_CENSUS), encoding="utf-8") as fh:
        rows = list(csv.reader(fh))
    start = next(i for i, r in enumerate(rows) if r[:3] == ["SA2 (UR)", "SEXP Sex", "AGEP Age"])
    for r in rows[start + 1:]:
        if len(r) < 4 or r[0] not in by_name:
            continue
        age = 115 if r[2].startswith("115") else int(r[2])
        sex = "M" if r[1] == "Male" else "F"
        out[(by_code[by_name[r[0]]], sex)][age] += int(r[3])
    return out


def read_erp(target):
    """ERP by (year, area, sex) -> five-year groups (0-4 .. 80-84, 85 and over)."""
    wb = openpyxl.load_workbook(os.path.join(target, K.F_ERP), read_only=True)
    out = {}
    for y in range(2016, 2027):
        ws = wb[str(y)]
        head = None
        for row in ws.iter_rows(values_only=True):
            if head is None:
                if row and row[0] == "LGA code":
                    head = list(row)
                continue
            if row[0] is None:
                continue
            sex = {"Males": "M", "Females": "F", "Persons": "P"}[row[2]]
            out[(y, row[1], sex)] = dict(zip(head[3:], row[3:]))
    wb.close()
    return out


def read_nsw(target):
    out = defaultdict(lambda: np.zeros(101))
    with open(os.path.join(target, K.F_NSW), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            age = 100 if r["Age"].startswith("100") else int(r["Age"])
            out[(int(r["Year"]), r["Sex"])][age] += int(r["Estimated resident population"])
    return out


# ------------------------------------------------------------------------------------------------- the life table
def life_e85(m, ax=0.5):
    """Life expectancy at 85 from rates at 85..94 and 95 and over (closed on its own rate)."""
    m = np.asarray(m, dtype=float)
    q = np.where(m > 0, m / (1 + (1 - ax) * m), 0.0)
    q = np.minimum(q, 1.0)
    q[-1] = 1.0
    l = np.cumprod(np.r_[1.0, 1 - q[:-1]])
    d = l * q
    L = l - (1 - ax) * d
    L[-1] = l[-1] / m[-1] if m[-1] > 0 else l[-1] * 2.0
    return float(L.sum())


def life_e85_reading(m, reading):
    """Life expectancy at 85 under other readings of how a rate becomes years lived: 'qm' takes the rate as the
    probability of dying in the year; 'exp' holds the rate constant through each year of age; 'halfclose' closes 95 and
    over with half a year for each survivor to 95."""
    m = np.asarray(m, dtype=float)
    if reading == "qm":
        q = np.minimum(m, 1.0)
    elif reading == "exp":
        q = 1 - np.exp(-m)
    else:
        q = np.minimum(np.where(m > 0, m / (1 + 0.5 * m), 0.0), 1.0)
    q[-1] = 1.0
    l = np.cumprod(np.r_[1.0, 1 - q[:-1]])
    d = l * q
    L = np.where(m > 0, d / np.where(m > 0, m, 1.0), l) if reading == "exp" else l - 0.5 * d
    L[-1] = 0.5 * l[-1] if reading == "halfclose" else l[-1] / m[-1]
    return float(L.sum())


def deaths_by_age(deaths, area, sexes=("F", "M")):
    D = np.zeros(NA)
    for (ar, sx, dob, dod, age, _st) in deaths:
        if ar == area and sx in sexes and dod.year in K.WINDOW and age >= 85:
            D[min(age, K.TOP) - 85] += 1
    return D


def erp85(E, area, y, sex="P"):
    return int(E[(y, area, sex)]["85 and over"])


# ------------------------------------------------------------------------------------------------- the reconstruction
def census_window(dob):
    """The census age of a person born on dob, and whether they fall in the ninth born 1 July to 10 August."""
    a = age_on(dob, CEN)
    late = (dob.month, dob.day) >= (7, 1) and (dob.month, dob.day) <= (8, 10)
    return a, late


def reconstruct(C, deaths, E, area, split=True, scale=True, by_sex=False, missing_dob="infer", floor=False):
    """Exposures at 85..94 and 95+ summed over 30 June 2024, 2025 and 2026 for one area.

    Each census age a (persons, or by sex) is followed as the people of its birth window; a share 41/365 of them were
    born 1 July to 10 August and are a year younger at 30 June than the rest. From each part go the deaths since
    census night of the people of that window. split=False puts the whole census age in the main part; scale=False
    leaves the closed cohort unscaled; by_sex rebuilds and scales each sex on its own; missing_dob: 'infer' places a
    death recorded without a birth date by its age and date of death, 'drop' leaves it out. floor=True gives a part
    of a census age that its deaths since exhaust no survivors instead of a negative count."""
    sexes = (("F",), ("M",)) if by_sex else (("F", "M"),)
    total = np.zeros(NA)
    for sx in sexes:
        cen = sum(C[(area, s)] for s in sx)
        phi = 41.0 / 365.0 if split else 0.0
        main = cen * (1 - phi)
        late = cen * phi
        # deaths since census night by census age and part, to each 30 June
        dmain = {y: np.zeros(116) for y in K.WINDOW}
        dlate = {y: np.zeros(116) for y in K.WINDOW}
        for (ar, s, dob, dod, age, _st) in deaths:
            if ar != area or s not in sx or dod <= CEN:
                continue
            if dob is None:
                if missing_dob == "drop":
                    continue
                # born between dod - (age+1) years and dod - age years: its middle day
                mid = dod - dt.timedelta(days=int((age + 0.5) * 365.25))
                a, lt = census_window(mid)
            else:
                a, lt = census_window(dob)
            if not split:
                lt = False
            for y in K.WINDOW:
                if dod <= dt.date(y, 6, 30) and 0 <= a < 116:
                    (dlate if lt else dmain)[y][a] += 1
        for y in K.WINDOW:
            k = y - 2021
            N = np.zeros(NA)
            for a in range(116):
                for part, cnt, dd, shift in (("main", main[a], dmain[y][a], k), ("late", late[a], dlate[y][a], k - 1)):
                    x = a + shift
                    if x < 85:
                        continue
                    N[min(x, K.TOP) - 85] += max(cnt - dd, 0.0) if floor else cnt - dd
            if scale:
                tgt = sum(erp85_sex(E, area, y, s) for s in sx)
                N = N * tgt / N.sum()
            total += N
    return total


def erp85_sex(E, area, y, sex):
    return int(E[(y, area, sex)]["85 and over"])


def shares_census(C, area):
    cen = C[(area, "F")] + C[(area, "M")]
    v = np.r_[cen[85:K.TOP], cen[K.TOP:].sum()]
    return v / v.sum()


def shares_shifted(C, area, y):
    cen = C[(area, "F")] + C[(area, "M")]
    k = y - 2021
    v = np.zeros(NA)
    for a in range(116):
        x = a + k
        if x >= 85:
            v[min(x, K.TOP) - 85] += cen[a]
    return v / v.sum()


def shares_nsw(N, y):
    v = N[(y, "F")] + N[(y, "M")]
    v = np.r_[v[85:K.TOP], v[K.TOP:].sum()]
    return v / v.sum()


def rungs(target):
    """Every rung's life expectancy at 85 by area, from the shipped files."""
    by_code, by_name = read_corr(target)
    deaths = read_deaths(target, by_code)
    C = read_census(target, by_name, by_code)
    E = read_erp(target)
    N = read_nsw(target)
    out = {}
    for area in K.AREA_NAMES:
        D = deaths_by_age(deaths, area)
        tot = {y: erp85(E, area, y) for y in K.WINDOW}
        r = {}
        r["C"] = sum(tot.values()) / D.sum()
        r["R1"] = life_e85(D / sum(shares_nsw(N, y) * tot[y] for y in K.WINDOW))
        r["R2"] = life_e85(D / sum(shares_census(C, area) * tot[y] for y in K.WINDOW))
        r["R2s"] = life_e85(D / sum(shares_shifted(C, area, y) * tot[y] for y in K.WINDOW))
        Eg = reconstruct(C, deaths, E, area)
        r["G"] = life_e85(D / Eg)
        r["G_nosplit"] = life_e85(D / reconstruct(C, deaths, E, area, split=False))
        r["G_closed"] = life_e85(D / reconstruct(C, deaths, E, area, scale=False))
        r["G_bysex"] = life_e85(D / reconstruct(C, deaths, E, area, by_sex=True))
        r["G_dropdob"] = life_e85(D / reconstruct(C, deaths, E, area, missing_dob="drop"))
        r["G_ax45"] = life_e85(D / Eg, ax=0.45)
        r["G_floor"] = life_e85(D / reconstruct(C, deaths, E, area, floor=True))
        r["D"] = D
        r["E"] = Eg
        r["tot"] = tot
        out[area] = r
    return out, deaths, C, E


# ------------------------------------------------------------------------------------------------- the after-hours asks
def read_ah(target):
    with open(os.path.join(target, K.F_AH), encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def _ts(s):
    return dt.datetime.fromisoformat(s[:19])


def ah_rows_85(rows, by_code):
    """Attendances in 2026 (arrival's local date) on residents of the twelve areas aged 85 and over at the time."""
    out = []
    for r in rows:
        area = by_code.get(r["residence_sa2_code"])
        if area is None or int(r["age_at_attendance"]) < 85:
            continue
        if not r["arrived_at"].startswith("2026-"):
            continue
        out.append((area, r))
    return out


def ah_golden(rows, by_code, gap_min=0):
    """Visits: the occasions the service attended a resident, a clinician who arrives on a case while the visit's
    clinicians are still there counted with them and a later return on the same case counted again (a new occasion
    when an arrival comes more than gap_min minutes after the previous occasion's last departure). Residents: one
    person per old key, the new system's external identifier joining a carried-over client's new key to their old
    one."""
    by_case = defaultdict(list)
    for r in rows:
        by_case[r["case_ref"]].append(r)
    occasions = []
    for case, rs in by_case.items():
        rs = sorted(rs, key=lambda r: r["arrived_at"])
        cur_end = None
        for r in rs:
            arr, dep = _ts(r["arrived_at"]), _ts(r["departed_at"])
            if cur_end is None or arr > cur_end + dt.timedelta(minutes=gap_min):
                occasions.append(r)
                cur_end = dep
            else:
                cur_end = max(cur_end, dep)
    keep = {id(r) for _a, r in ah_rows_85(occasions, by_code)}
    visits = defaultdict(int)
    for r in occasions:
        if id(r) in keep:
            visits[by_code[r["residence_sa2_code"]]] += 1
    people = defaultdict(set)
    for area, r in ah_rows_85(rows, by_code):
        people[area].add(r["external_id"] or r["client_key"])
    return ({a: visits.get(a, 0) for a in K.AREA_NAMES}, {a: len(people.get(a, ())) for a in K.AREA_NAMES})


def ah_paths(rows, by_code):
    """The routes a careful solver may take through the attendance file."""
    sel = ah_rows_85(rows, by_code)
    out = {}
    rows_n = defaultdict(int)
    cases = defaultdict(set)
    keys = defaultdict(set)
    attrs = defaultdict(set)
    day_client = defaultdict(set)
    for area, r in sel:
        rows_n[area] += 1
        cases[area].add(r["case_ref"])
        keys[area].add(r["client_key"])
        attrs[area].add((r["birth_month"], r["sex"], r["residence_sa2_code"]))
        day_client[area].add((r["client_key"], r["arrived_at"][:10]))
    out["visits_rows"] = {a: rows_n.get(a, 0) for a in K.AREA_NAMES}
    out["visits_cases"] = {a: len(cases.get(a, ())) for a in K.AREA_NAMES}
    out["visits_client_day"] = {a: len(day_client.get(a, ())) for a in K.AREA_NAMES}
    out["residents_keys"] = {a: len(keys.get(a, ())) for a in K.AREA_NAMES}
    out["residents_attrs"] = {a: len(attrs.get(a, ())) for a in K.AREA_NAMES}
    return out
