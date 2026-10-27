"""task270 records the call reads: the mood app's two exports (version 3 to 14 June 2026, version 4 from 15 June),
the enrolment log, the assessment fortnights register, the care episodes, and the questionnaire scores taken at each
assessment (a file the call does not read). Each file carries its own conventions the way its system writes them."""
from __future__ import annotations

import csv
import datetime as dt
import os

import numpy as np
import openpyxl
from openpyxl.styles import Font

import common as K


def _w(path, header, rows):
    with open(path, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh, lineterminator="\r\n")
        w.writerow(header)
        w.writerows(rows)
    return len(rows)


def items_v3(vals):
    return [int(v) for v in np.clip(np.round(vals * 10), 0, 100)]


def items_v4(vals):
    return [int(v) for v in np.clip(np.round(vals), 0, 10)]


def positive(r, vals, mu):
    na = float(np.mean(vals))
    cheer = 6.4 - 0.55 * (na - mu) - 0.35 * (mu - 4.5) + r.normal(0, 1.1)
    relax = 5.9 - 0.5 * (na - mu) - 0.3 * (mu - 4.5) + r.normal(0, 1.2)
    return np.array([cheer, relax])


def prompt_rows(patients, extra):
    """Every prompt as one record: (code, patient, sent local, answered local, block, NA item values, PA values,
    company)."""
    out = []
    for p in patients + extra:
        rp = K.rng("export", p.number)
        for kind, first, c, swap, c2, rows in p.fortnight_rows:
            for k, b, sent, answered, vals in rows:
                cc = c2 if (swap is not None and k >= swap) else c
                pa = positive(rp, vals, p.mu) if vals is not None else None
                company = None
                if vals is not None:
                    company = "with others" if rp.random() < (0.35 + 0.25 * (sent.hour >= 17)) else "alone"
                out.append((cc, p.number, sent, answered, b, vals, pa, company))
    out.sort(key=lambda z: (z[2], z[0]))
    return out


def write_exports(out_dir, prompts):
    n3 = n4 = 0
    v3, v4 = [], []
    r = K.rng("v4 ids")
    used = set()
    resend = K.rng("resends")
    for cc, num, sent, answered, b, vals, pa, company in prompts:
        if sent.date() < K.V4_FROM:
            if vals is None:
                row = [cc, b + 1, sent.strftime("%d/%m/%Y %H:%M"), "", "", "", "", "", "", "", ""]
            else:
                row = [cc, b + 1, sent.strftime("%d/%m/%Y %H:%M"), answered.strftime("%d/%m/%Y %H:%M:%S")] + \
                      items_v3(vals) + [int(v) for v in np.clip(np.round(pa * 10), 0, 100)] + [company]
            v3.append(row)
        else:
            while True:
                pid = "%010d" % int(r.integers(1_000_000_000, 9_999_999_999))
                if pid not in used:
                    used.add(pid)
                    break
            s_utc = K.local_to_utc(sent)
            build = "4.0.3" if sent.date() < dt.date(2026, 8, 24) else "4.1.0"
            if vals is None:
                row = [cc, pid, b + 1, s_utc.strftime("%Y-%m-%dT%H:%M:%SZ"), "", "expired", "", "", "", "", "", "",
                       "", build]
                v4.append(row)
            else:
                a_utc = K.local_to_utc(answered)
                row = [cc, pid, b + 1, s_utc.strftime("%Y-%m-%dT%H:%M:%SZ"), a_utc.strftime("%Y-%m-%dT%H:%M:%SZ"),
                       "answered"] + items_v4(vals) + [int(v) for v in np.clip(np.round(pa), 0, 10)] + \
                      ["Y" if company == "with others" else "N", build]
                v4.append(row)
                if resend.random() < 0.021:
                    v4.append(list(row))
    n3 = _w(os.path.join(out_dir, K.F_V3), ["participant_id", "beep", "scheduled_local", "responded_local",
                                            "na_down_vas", "na_anxious_vas", "na_irritable_vas", "na_ashamed_vas",
                                            "pa_cheerful_vas", "pa_relaxed_vas", "company"], v3)
    n4 = _w(os.path.join(out_dir, K.F_V4), ["participant", "prompt_id", "block", "sent_at_utc", "answered_at_utc",
                                            "status", "down", "anxious", "irritable", "ashamed", "cheerful", "relaxed",
                                            "with_others", "app_build"], v4)
    return {K.F_V3: n3, K.F_V4: n4}


def write_enrolments(out_dir, patients, extra):
    rows = []
    for p in patients + extra:
        for kind, first, c, swap, c2, _rows in p.fortnight_rows:
            enrol_on = [e for e in p.codes if e[0] == c][0][1]
            last = first + dt.timedelta(days=13)
            if swap is not None:
                rows.append([c, p.number, enrol_on.isoformat(), (first + dt.timedelta(days=swap - 1)).isoformat(),
                             "replaced phone"])
                rows.append([c2, p.number, (first + dt.timedelta(days=swap - 1)).isoformat(), last.isoformat(),
                             "assessment ended"])
            elif kind in ("ended_early", "incomplete"):
                rows.append([c, p.number, enrol_on.isoformat(), last.isoformat(),
                             "withdrew" if kind == "incomplete" else "assessment ended"])
            else:
                rows.append([c, p.number, enrol_on.isoformat(), last.isoformat(), "assessment ended"])
    rows.sort(key=lambda z: (z[2], z[0]))
    return {K.F_ENROL: _w(os.path.join(out_dir, K.F_ENROL),
                          ["participant_code", "patient_number", "enrolled_on", "ended_on", "end_reason"], rows)}


def team_on(p, day):
    for team, s, e, _ in p.episodes:
        if s <= day and (e is None or e >= day):
            return team
    return None


def write_register(out_dir, patients, extra):
    r = K.rng("register")
    used = set()
    recs = []
    for p in patients + extra:
        for kind, first, c, swap, c2, rows in p.fortnight_rows:
            answered = sum(1 for z in rows if z[3] is not None)
            if kind in ("latest", "earlier") and answered >= K.COMPLETE:
                status = "Completed"
            elif kind == "incomplete":
                status = "Completed" if r.random() < 0.28 else ("Withdrawn" if r.random() < 0.6 else "Ended early")
            else:
                status = "Completed" if r.random() < 0.25 else "Ended early"
            team = team_on(p, first)
            req = team
            if r.random() < 0.03:
                req = [c0 for c0, _ in K.TEAMS if c0 != K.CURRENT[team]][int(r.integers(0, 10))]
            while True:
                ref = "AF26-%04d" % int(r.integers(1000, 9999))
                if ref not in used:
                    used.add(ref)
                    break
            recs.append([ref, int(p.number), c, first, first + dt.timedelta(days=13),
                         K.TEAM_NAME[req] + " CMHT", status])
    recs.sort(key=lambda z: (z[3], z[0]))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fortnights 2026"
    head = ["Ref", "Patient number", "App code", "First day", "Last day", "Requested by", "Status"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    for rec in recs:
        ws.append(rec)
    for row in ws.iter_rows(min_row=2):
        row[3].number_format = row[4].number_format = "dd/mm/yyyy"
    for col, w in zip("ABCDEFG", (11, 15, 12, 12, 12, 26, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(os.path.join(out_dir, K.F_REGISTER))
    return {K.F_REGISTER: len(recs)}


def write_episodes(out_dir, patients, extra, hub_patients):
    rows = []
    seen = set()
    for p in patients + extra + hub_patients:
        if p.number in seen:
            continue
        seen.add(p.number)
        for team, s, e, why in p.episodes:
            rows.append([p.number, team, K.TEAM_NAME[team] + " CMHT", s.isoformat(), e.isoformat() if e else "", why])
    rows.sort(key=lambda z: (z[0], z[3]))
    return {K.F_EPISODES: _w(os.path.join(out_dir, K.F_EPISODES),
                             ["patient_number", "team_code", "team_name", "episode_start", "episode_end", "end_reason"],
                             rows)}


def write_phq(out_dir, patients, extra):
    rows = []
    for p in patients + extra:
        r = K.rng("phq", p.number)
        for kind, first, c, swap, c2, _rows in p.fortnight_rows:
            phq = int(np.clip(round(1.9 * p.mu + 2.0 + r.normal(0, 3.4)), 0, 27))
            gad = int(np.clip(round(1.5 * p.mu + 1.4 + r.normal(0, 3.0)), 0, 21))
            rows.append([p.number, first.strftime("%d/%m/%Y"), phq, gad])
    rows.sort(key=lambda z: (z[1][6:] + z[1][3:5] + z[1][:2], z[0]))
    return {K.F_PHQ: _w(os.path.join(out_dir, K.F_PHQ), ["patient_number", "assessment_date", "phq9_total",
                                                         "gad7_total"], rows)}
