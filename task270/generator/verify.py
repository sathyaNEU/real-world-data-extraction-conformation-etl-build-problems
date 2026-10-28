"""task270 independent verifier: reads only target/, the prompt and the guidelines' filed terms, recomputes every
graded figure by its own route and compares it with figures.json and submission.md. It shares no code with calc.py.

    python3 task270/generator/verify.py"""
from __future__ import annotations

import csv
import datetime as dt
import json
import os
import re
import sys
from collections import defaultdict

import numpy as np
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
T = os.path.join(TASK, "target")
CENSUS_NIGHT = dt.date(2021, 8, 10)
YEARS = (2024, 2025, 2026)
OPEN = 95          # guidelines 3.2: single years 85 to 94, 95 and over last
LINE = 7.0         # guidelines 3.1


def find(prefix):
    return next(os.path.join(T, f) for f in os.listdir(T) if f.startswith(prefix))


def completed_years(born, on):
    return on.year - born.year - ((on.month, on.day) < (born.month, born.day))


def load():
    sa2_area, name_code = {}, {}
    for r in csv.DictReader(open(find("asgs2021_sa2_to_lga2023"), encoding="utf-8")):
        sa2_area[r["SA2_CODE_2021"]] = r["LGA_NAME_2023"]
        name_code[r["SA2_NAME_2021"]] = r["SA2_CODE_2021"]
    deaths = []
    for r in csv.DictReader(open(find("bdm_death_registrations"), encoding="utf-8")):
        a = sa2_area.get(r["usual_residence_sa2_code"])
        if a:
            deaths.append([a, r["sex"], r["date_of_birth"] or None, r["date_of_death"], int(r["age_at_death"])])
    wb = openpyxl.load_workbook(find("interstate_deaths"), read_only=True)
    rows = list(wb["Deaths"].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "State of registration")
    for r in rows[hi + 1:]:
        if r[0] is None:
            continue
        a = sa2_area.get(str(r[7]))
        if a:
            deaths.append([a, r[5], r[4].date().isoformat(), r[3].date().isoformat(), int(r[6])])
    wb.close()
    census = defaultdict(lambda: np.zeros(116))
    lines = list(csv.reader(open(find("census2021"), encoding="utf-8")))
    h = next(i for i, r in enumerate(lines) if r[:1] == ["SA2 (UR)"])
    for r in lines[h + 1:]:
        if len(r) >= 4 and r[0] in name_code and r[3].strip().isdigit():
            age = 115 if r[2].startswith("115") else int(r[2])
            census[sa2_area[name_code[r[0]]]][age] += int(r[3])
    erp = {}
    wb = openpyxl.load_workbook(find("erp_lga_age_sex"), read_only=True)
    for y in YEARS:
        rows = list(wb[str(y)].iter_rows(values_only=True))
        hi = next(i for i, r in enumerate(rows) if r and r[0] == "LGA code")
        col = list(rows[hi]).index("85 and over")
        for r in rows[hi + 1:]:
            if r[0] and r[2] == "Persons":
                erp[(r[1], y)] = int(r[col])
    wb.close()
    return sa2_area, deaths, census, erp


def life_expectancy_85(rates):
    """Single-year table from 85 with an open last group; half a year lived in the year of death."""
    l, total = 1.0, 0.0
    for i, m in enumerate(rates):
        if i == len(rates) - 1:
            total += l / m
            break
        q = min(m / (1 + 0.5 * m), 1.0)
        total += l * (1 - 0.5 * q)
        l *= 1 - q
    return total


def e85(area, deaths, census, erp):
    # deaths at each age, 2024 to 2026 by date of death
    D = np.zeros(OPEN - 84)
    for a, s, dob, dod, age in deaths:
        if a == area and int(dod[:4]) in YEARS and age >= 85:
            D[min(age, OPEN) - 85] += 1
    # each census age's birth window: born 11 August (2020 - c) to 10 August (2021 - c); the part born 1 July to
    # 10 August is a year younger at 30 June than the rest
    late_share = 41 / 365
    cen = census[area]
    gone = {y: defaultdict(float) for y in YEARS}      # (census age, late part?) -> deaths since census night
    for a, s, dob, dod, age in deaths:
        if a != area:
            continue
        dd = dt.date.fromisoformat(dod)
        if dd <= CENSUS_NIGHT:
            continue
        if dob:
            b = dt.date.fromisoformat(dob)
        else:
            b = dd - dt.timedelta(days=int((age + 0.5) * 365.25))
        c = completed_years(b, CENSUS_NIGHT)
        late = (7, 1) <= (b.month, b.day) <= (8, 10)
        for y in YEARS:
            if dd <= dt.date(y, 6, 30):
                gone[y][(c, late)] += 1
    expo = np.zeros(OPEN - 84)
    for y in YEARS:
        pop = np.zeros(OPEN - 84)
        for c in range(116):
            for late, share in ((False, 1 - late_share), (True, late_share)):
                x = c + (y - 2021) - (1 if late else 0)
                if x >= 85:
                    pop[min(x, OPEN) - 85] += cen[c] * share - gone[y][(c, late)]
        expo += pop * erp[(area, y)] / pop.sum()
    return life_expectancy_85(D / expo)


def asks(sa2_area):
    rows = list(csv.DictReader(open(find("afterhours_attendances"), encoding="utf-8")))
    visits, people = defaultdict(int), defaultdict(set)
    cases = defaultdict(list)
    for r in rows:
        cases[r["case_ref"]].append(r)
    for case, rs in cases.items():
        rs.sort(key=lambda r: r["arrived_at"])
        last_dep = None
        for r in rs:
            arr = dt.datetime.fromisoformat(r["arrived_at"][:19])
            dep = dt.datetime.fromisoformat(r["departed_at"][:19])
            new = last_dep is None or arr > last_dep      # a clinician arriving while the visit is on joins it
            last_dep = dep if new else max(last_dep, dep)
            a = sa2_area.get(r["residence_sa2_code"])
            if new and a and int(r["age_at_attendance"]) >= 85 and r["arrived_at"][:4] == "2026":
                visits[a] += 1
    for r in rows:
        a = sa2_area.get(r["residence_sa2_code"])
        if a and int(r["age_at_attendance"]) >= 85 and r["arrived_at"][:4] == "2026":
            people[a].add(r["external_id"] or r["client_key"])
    return dict(visits), {a: len(v) for a, v in people.items()}


def main():
    sa2_area, deaths, census, erp = load()
    areas = sorted(set(sa2_area.values()))
    F = json.load(open(os.path.join(HERE, "figures.json")))
    sub = open(os.path.join(TASK, "submission.md")).read()
    fails = []
    e = {a: e85(a, deaths, census, erp) for a in areas}
    call = sorted(a for a in areas if e[a] < LINE)
    if call != F["call"]:
        fails.append(("call", call, F["call"]))
    for a in areas:
        if round(e[a], 1) != F["e85"][a]:
            fails.append(("e85", a, round(e[a], 3), F["e85"][a]))
        if not re.search(r"- %s: %.1f, (qualifies|does not qualify)" % (re.escape(a), F["e85"][a]), sub):
            fails.append(("submission e85", a))
    v, p = asks(sa2_area)
    for a in areas:
        if v.get(a, 0) != F["visits"][a]:
            fails.append(("visits", a, v.get(a, 0), F["visits"][a]))
        if p.get(a, 0) != F["residents"][a]:
            fails.append(("residents", a, p.get(a, 0), F["residents"][a]))
    prompt = open(os.path.join(TASK, "prompt.txt")).read()
    if "seven" not in prompt or "one decimal place" not in prompt:
        fails.append(("prompt", "line or rounding"))
    for f in fails:
        print("FAIL", f)
    print("verify: %d areas, call %s, %d failed" % (len(areas), ", ".join(call), len(fails)))
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
