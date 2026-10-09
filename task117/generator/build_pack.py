"""task117 pack build.

    python3 task117/generator/build_pack.py --out <dir for target files> --meta <path of metadata.json>
                                            [--record <path for the build record json>]

Seeded and deterministic: two runs write byte-identical files. Every assertion in checks.py runs
before a file is written; the pack gates run after."""
from __future__ import annotations

import argparse
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import date, datetime

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, HERE)

import checks as C  # noqa: E402
import documents as D  # noqa: E402
import pipeline  # noqa: E402
import writers as Wr  # noqa: E402
from analysis import GROWTH, Analysis  # noqa: E402
from common import iso_local, nearest5  # noqa: E402

SCRUB = os.path.join(REPO, ".claude", "skills", "reduce-house-fixes", "scripts", "scrub_producer_metadata.py")
EXPORT_TS = datetime(2027, 1, 18, 9, 0, 0)
FLOOR, CEILING = "2023-01-01", "2027-01-25"

# in-fiction author and date of every binary input
PRODUCERS = {
    Wr.LOG: ("City of Larch Harbor Facilities", "2027-01-04", (2027, 1, 4, 8, 12, 0)),
    Wr.AGREEMENT: ("North Sound Power & Light", "2027-01-12", (2027, 1, 12, 15, 40, 0)),
    Wr.RATES: ("North Sound Power & Light", "2025-11-14", None),
    Wr.STANDARD: ("City of Larch Harbor", "2026-09-15", None),
    Wr.GUIDE: ("North Sound Power & Light", "2026-03-02", None),
}

SOURCES = [
    (Wr.HEADER, "Curbline Charging settlement export, all eight city garages: sessions plugged in from January 1, "
                "2024 through December 31, 2026, every weekly delivery received through January 18, 2027, kept as "
                "delivered."),
    (Wr.SPINE, "Curbline interval export for the same session records, one row per quarter-hour from plug-in to "
               "plug-out."),
    (Wr.GATEWAY, "Sessions at the units that reported through the legacy gateway B until it was retired on April 22, "
                 "2024. These sessions are not in the settlement export. Produced by Curbline from the gateway "
                 "archive, May 2, 2024."),
    (Wr.DECISIONS, "Parking Services decisions on restated session versions received from Curbline in 2025. A "
                   "restated version marked ACCEPTED replaces the version before it; a version marked REJECTED "
                   "leaves the earlier version in force."),
    (Wr.REGISTER, "Curbline station register: every station identifier with its garage, unit position, model, "
                  "rating and in-service dates. Extract of January 18, 2027."),
    (Wr.PERMITS, "Parking Services permit system: Civic Center EV charging permits, all statuses, as of January 18, "
                 "2027."),
    (Wr.CHECKS, "Parking Services vehicle checks against registration at permit issue and renewal, 2021 to January "
                "15, 2027."),
    (Wr.REFERENCE, "Parking Services vehicle reference table, compiled from manufacturer specification sheets. "
                   "Body class sets the permit fee class. Updated November 2, 2026."),
    (Wr.FLEET, "Fleet Services roster of city vehicles that hold fleet cards, as of January 11, 2027."),
    (Wr.FLEETCARD, "Fleet Services: EV charging transactions on city fleet cards at Curbline stations, January 2024 to "
                   "December 2026, from the fleet card processor's monthly statement files. NETWORK_REF is the "
                   "charging network's reference for the charge."),
    (Wr.COURTESY, "Curbline courtesy-session report for the Civic Center decks, January 2024 to December 2026: "
                  "charging on weekends and Schedule 26 holidays, when it is free to permit holders."),
    (Wr.STATEMENTS, "Parking Services monthly EV charging statements to Civic Center permit holders, 2026: per permit "
                    "and month, the charges billed (a session fee of $0.50 each) and the energy (21.8 cents per kWh)."),
    (Wr.LOG, "Facilities Electrical: monthly reads of the two deck charging panel sub-meters, December 2023 to "
             "December 2026, as kept by the electricians."),
    (Wr.NAMEPLATES, "Facilities Electrical: nameplate and configuration record of the two deck sub-meters."),
    (Wr.SCHEDULE, "Facilities Electrical: panel schedules for the deck charging panels and the North Deck house "
                  "panel, with the dates each circuit assignment applied."),
    (Wr.WORKORDERS, "Facilities work order system, 2026, Civic Center buildings."),
    (Wr.CAMPUS, "North Sound Power & Light statements for the Civic Center campus account, 2024 to 2026, keyed "
                "from the paper bills by Finance."),
    (Wr.STATUS, "Curbline station status events for all eight garages, 2026."),
    (Wr.RATES, "North Sound Power & Light tariff, Schedule 26, as issued November 14, 2025."),
    (Wr.GUIDE, "North Sound Power & Light New Service Planning Guide, 2026 edition, Section 7, as provided by the "
               "NSPL planner."),
    (Wr.AGREEMENT, "Draft service agreement from NSPL, January 12, 2027, for the Council packet."),
    (Wr.STANDARD, "City standard FES-07, Revision 4."),
    (Wr.NOTES, "Curbline's field notes for the settlement export."),
]


def normalise_zip(path, dt):
    """Rewrite an OOXML container with fixed entry timestamps and order, contents unchanged."""
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        data = [(i.filename, z.read(i.filename)) for i in infos]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for name, blob in data:
            zi = zipfile.ZipInfo(name, date_time=dt)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            out.writestr(zi, blob)
    with open(path, "wb") as f:
        f.write(buf.getvalue())


def scrub(path, producer, stamp):
    if path.endswith((".docx", ".xlsx")):
        producer = producer.replace("&", "&amp;")  # written into OOXML core properties as XML text
    with tempfile.TemporaryDirectory() as td:
        tmp = os.path.join(td, os.path.basename(path))
        shutil.copy2(path, tmp)
        r = subprocess.run([sys.executable, SCRUB, td, "--apply", "--producer", producer, "--stamp", stamp,
                            "--floor", FLOOR, "--ceiling", CEILING], capture_output=True, text=True)
        a = subprocess.run([sys.executable, SCRUB, td, "--floor", FLOOR, "--ceiling", CEILING], capture_output=True,
                           text=True)
        C.ck(f"H1 container scrubbed: {os.path.basename(path)}", a.returncode == 0 and "clean" in a.stdout,
             r.stdout + r.stderr + a.stdout)
        shutil.copyfile(tmp, path)


def doc_texts(out):
    """Text of every document in the pack, for the signpost and leak greps."""
    from pypdf import PdfReader
    texts = {}
    for f in sorted(os.listdir(out)):
        p = os.path.join(out, f)
        if f.endswith(".pdf"):
            texts[f] = "\n".join(pg.extract_text() for pg in PdfReader(p).pages)
        elif f.endswith(".docx"):
            from docx import Document
            d = Document(p)
            parts = [x.text for x in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.extend(c.text for c in row.cells)
            texts[f] = "\n".join(parts)
        elif f.endswith(".txt"):
            texts[f] = open(p).read()
        elif f.endswith(".xlsx"):
            from openpyxl import load_workbook
            wb = load_workbook(p, read_only=True)
            parts = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    parts.extend(str(x) for x in row if isinstance(x, str))
            texts[f] = "\n".join(parts)
    return texts


SIGNPOST = [r"\bon-?board\b", r"\bvehicle'?s? (charger|charging (rate|speed|power))", r"\baccepts?\b",
            r"\b(limit|limits|limited|capped)\b[^.]{0,50}\b(draw|charging power|charging rate)",
            r"\bdraws?\b[^.]{0,60}\b(vehicle|car)\b", r"\b(vehicle|car)\b[^.]{0,60}\bdraws?\b"]


def pack_gates(out, meta, prompt_path, answer_figures):
    files = sorted(os.listdir(out))
    fmts = sorted({os.path.splitext(f)[1] for f in files})
    C.ck("G01 input gate: 10 or more files", len(files) >= 10, len(files))
    C.ck("G02 input gate: 3 or more formats", len(fmts) >= 3, fmts)
    import pyarrow.parquet as pq
    n = pq.ParquetFile(os.path.join(out, Wr.SPINE)).metadata.num_rows
    C.ck("G03 input gate: a file of 25,000 or more rows", n >= 25000, n)
    C.ck("G04 input gate: two or more distractors named in metadata.json and present in the pack",
         len(meta["distractor_files"]) >= 2 and all(d in files for d in meta["distractor_files"]))
    leaks = []
    for f in files:
        p = os.path.join(out, f)
        blob = open(p, "rb").read()
        if re.search(rb"(?i)distractor", blob) or re.search(r"(?i)distractor", f):
            leaks.append(f)
    C.ck("G05 the word distractor appears nowhere under target/", not leaks, leaks)
    texts = doc_texts(out)
    hits = []
    for f, t in texts.items():
        flat = re.sub(r"\s+", " ", t)
        for pat in SIGNPOST:
            for m in re.finditer(pat, flat, flags=re.I):
                hits.append((f, m.group(0)))
    C.ck("G06 anti-signpost grep: no document says what limits a session's draw or ties charging power to the vehicle",
         not hits, hits)
    fig_hits = [(f, x) for f, t in texts.items() for x in answer_figures if re.search(r"(?<![\d.])" + re.escape(x)
                                                                                    + r"(?![\d])", t)]
    C.ck("G07 no golden figure appears in any document", not fig_hits, fig_hits)
    em = [f for f, t in texts.items() if "\u2014" in t] + [f for f in files if f.endswith(".csv") and
                                                          "\u2014" in open(os.path.join(out, f)).read()]
    C.ck("G08 no em dash anywhere in the pack", not em, em)
    prompt = open(prompt_path).read()
    need = ["service agreement", "contracted demand", "North Sound Power & Light", "I buy power"]
    banned = ["billing hours", "onboard", "accept"]
    C.ck("G09 H20: the prompt carries the sourcing nouns and none of the stump's words",
         all(x in prompt for x in need) and not any(x in prompt.lower() for x in banned))
    C.ck("G10 the deliverables the prompt names are not in the pack",
         not any(x in files for x in ("contract_demand_note.pdf", "deck_load_day.png", "civic_service_demand.xlsx")))
    # single-statement invariant: each load-bearing rule stated in exactly one shipped file
    rules = {"billing window": r"12:00 noon through 7:45 p\.?m\.?", "growth factor 1.12": r"(?<![,\d.])1\.12(?![\d,])",
             "latest twelve closed months": r"latest twelve closed calendar months",
             "one revenue meter": r"one NSPL revenue meter", "interval start convention": r"Start of the quarter-hour",
             "auth code persistence": r"keeps its code", "later reading stands": r"later reading stands",
             "meter on standard time": r"Pacific Standard Time", "nameplate sizing": r"Table 7-2",
             "accuracy record": r"recorded billing demand in whole kilowatts",
             "version of record": r"marked ACCEPTED replaces", "gateway B coverage": r"not in the settlement export",
             "permit-only decks": r"permit-only", "settlement run time": r"settlement run is daily",
             "free courtesy charging": r"free to permit holders"}
    alltext = dict(texts)
    for f in files:
        if f.endswith(".csv"):
            alltext[f] = open(os.path.join(out, f)).read()
    counts = {k: sorted(f for f, t in alltext.items() if re.search(p, re.sub(r"\s+", " ", t))) for k, p in rules.items()}
    C.ck("G11 single-statement invariant: each load-bearing rule stated in exactly one file",
         all(len(v) == 1 for v in counts.values()), counts)
    rows = {f: sum(1 for _ in open(os.path.join(out, f))) - 1 for f in files if f.endswith(".csv")}
    C.ck("G13 generation tell: no two data files share a row count", len(set(rows.values())) == len(rows), rows)
    return {"files": files, "formats": fmts, "spine_rows": n, "rule_homes": counts, "csv_rows": rows}


def write_all(w, out, rng):
    os.makedirs(out, exist_ok=True)
    for f in os.listdir(out):
        os.remove(os.path.join(out, f))
    j = lambda f: os.path.join(out, f)  # noqa: E731
    info = {}
    info["header_rows"] = len(Wr.write_header(w, j(Wr.HEADER)))
    info["spine_rows"] = Wr.write_spine(w, j(Wr.SPINE))
    Wr.write_decisions(w, j(Wr.DECISIONS), rng)
    Wr.write_gateway(w, j(Wr.GATEWAY), rng)
    Wr.write_register(w, j(Wr.REGISTER))
    Wr.write_schedule(j(Wr.SCHEDULE))
    Wr.write_log(w, j(Wr.LOG))
    Wr.write_nameplates(j(Wr.NAMEPLATES))
    Wr.write_work_orders(j(Wr.WORKORDERS))
    Wr.write_permits(w, j(Wr.PERMITS))
    Wr.write_checks(w, j(Wr.CHECKS))
    Wr.write_reference(j(Wr.REFERENCE))
    Wr.write_fleet(w, j(Wr.FLEET))
    Wr.write_fleet_card(w, j(Wr.FLEETCARD), rng)
    Wr.write_courtesy(w, j(Wr.COURTESY), rng)
    Wr.write_statements(w, j(Wr.STATEMENTS))
    Wr.write_campus(w, j(Wr.CAMPUS), rng)
    Wr.write_status(w, j(Wr.STATUS), rng)
    D.rate_schedule(j(Wr.RATES))
    D.forecasting_standard(j(Wr.STANDARD))
    D.planning_guide(j(Wr.GUIDE))
    D.service_agreement(j(Wr.AGREEMENT))
    D.field_notes(j(Wr.NOTES))
    D.data_sources(j(Wr.SOURCES), SOURCES)
    for f, (producer, stamp, zdt) in PRODUCERS.items():
        if zdt:
            normalise_zip(j(f), zdt)
        scrub(j(f), producer, stamp)
    ts = EXPORT_TS.timestamp()
    for f in os.listdir(out):
        os.utime(j(f), (ts, ts))
    return info


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    ap.add_argument("--meta", required=True)
    ap.add_argument("--record", default=None)
    ap.add_argument("--prompt", default=os.path.join(HERE, "..", "prompt.md"))
    a = ap.parse_args()

    w = pipeline.build_world()
    an = Analysis(w)
    r = an.ladder()
    main_res = C.main_call(an, r)
    renewal_res = C.renewal(an)
    charges_res = C.charges(an)
    corpus_res = C.corpus(an, r)
    b3_res = C.b3(an)
    b1_res = C.b1(an)
    counts, repairs = C.separation(an, r, main_res["answer"])
    C.referee(an)

    rng = np.random.default_rng(pipeline.SEED + 99)
    info = write_all(w, a.out, rng)
    meta = {
        "task": "task117",
        "domain": "Supply Chain & Logistics",
        "subdomain": "sourcing-procurement",
        "objective": "Forecasting & Predictive Modeling",
        "as_of": "2027-01-25",
        "deliverables": ["contract_demand_note.pdf", "deck_load_day.png", "civic_service_demand.xlsx"],
        "distractor_files": list(Wr.DISTRACTORS),
        "source": ("Every file is constructed for this task around a fictional city (Larch Harbor, Washington), "
                   "utility (North Sound Power & Light), charging network (Curbline Charging) and county. No real "
                   "person, account or record is depicted. Vehicle makes, models and onboard charger ratings in the "
                   "reference list follow the manufacturers' published specifications."),
        "license": "CC BY 4.0",
        "created": "2026-10-09",
        "generator": "task117/generator/build_pack.py (seeded, deterministic)",
        "files": [],
    }
    for f in sorted(os.listdir(a.out)):
        p = os.path.join(a.out, f)
        meta["files"].append({"path": f, "format": os.path.splitext(f)[1][1:], "bytes": os.path.getsize(p)})
    figs = ["129.136", "129.14", "115.3", "82.88", "46.256", "103.04", "98.56", "412.16", "309.12", "148.512",
            "132.6"]
    gates = pack_gates(a.out, meta, os.path.abspath(a.prompt), figs)
    heads = [main_res["answer"], main_res["rung0"], main_res["rung1"], main_res["rung2"], main_res["rung3"],
             main_res["rung4"], main_res["rung4_rec"], main_res["rung5_rec"]]
    C.ck("G14 generation tell: no headline figure sits on a round boundary",
         all(abs(x / 5 - round(x / 5)) > 0.02 for x in heads), heads)
    with open(a.meta, "w") as f:
        json.dump(meta, f, indent=2)
        f.write("\n")
    C.ck("G12 metadata.json names the distractors and carries no answer figure",
         all(x not in json.dumps(meta) for x in figs) and meta["distractor_files"])

    record = {
        "answer_unrounded": round(main_res["answer"], 3), "answer_filed": nearest5(main_res["answer"]),
        "split": [round(x, 3) for x in main_res["split"]], "split_unscaled": [round(x, 3) for x in main_res["split_unscaled"]],
        "monthly": [round(x, 3) for x in main_res["monthly"]],
        "rungs": {"0": round(main_res["rung0"], 3), "1": round(main_res["rung1"], 3), "2": round(main_res["rung2"], 3),
                  "3": round(main_res["rung3"], 3), "4": round(main_res["rung4"], 3)},
        "rung4_split": [round(x, 3) for x in main_res["rung4_split"]],
        "records": {"rung4": round(main_res["rung4_rec"], 3), "rung5": round(main_res["rung5_rec"], 3),
                    "rung5_at": main_res["rung5_rec_at"],
                    "rung5_split": [round(x, 3) for x in main_res["rung5_rec_split"]],
                    "rung5_monthly": [round(x, 3) for x in main_res["rung5_rec_monthly"]],
                    "rung4_monthly": [round(x, 3) for x in main_res["rung4_rec_monthly"]]},
        "charges": charges_res,
        "rung4_monthly": [round(x, 3) for x in main_res["rung4_monthly"]],
        "renewal": {k: v for k, v in renewal_res.items()},
        "cells": {k: round(v, 3) for k, v in main_res["cells"].items()},
        "two_error": {k: round(v, 3) for k, v in main_res["two_error"].items()},
        "convergent": {k: round(v, 3) for k, v in main_res["convergent"].items()},
        "family": corpus_res["family"], "twins_ks": corpus_res["twins_ks"], "shares": corpus_res["shares"],
        "departure_min_slack_min": corpus_res["departure_min_slack_min"],
        "b3_forecast": b3_res["forecast"], "b3_miss": b3_res["miss"],
        "b3_subset_mean_range": b3_res["subset_mean_range"],
        "b1_golden": {f"{p} {k:02d}": v for (p, k), v in b1_res["golden"].items()},
        "b1_clock_moves": {f"{p} {k:02d}": v for (p, k), v in b1_res["clock_moves"].items()},
        "b1_redelivery_moves": b1_res["redelivery_moves"], "b1_backfeed_moves": b1_res["backfeed_moves"],
        "b1_double_read_move": b1_res["double_read_move"], "b1_nearest_wrong": b1_res["nearest_wrong"],
        "b1_courtesy_moves": {f"{p} {k:02d}": v for (p, k), v in b1_res["courtesy_moves"].items()},
        "b1_split_moves": {f"{p} {k:02d}": v for (p, k), v in b1_res["split_moves"].items()},
        "b3_moves": b3_res["moves"],
        "separation_counts": counts, "pack": {k: v for k, v in gates.items() if k != "rule_homes"},
        "rule_homes": gates["rule_homes"], "assertions": len(C.LOG),
    }
    for name, ok, _ in C.LOG:
        print(("PASS " if ok else "FAIL ") + name)
    print(f"\n{len(C.LOG)} assertions, all green. Answer {record['answer_unrounded']} kW, files "
          f"{record['answer_filed']} kW.")
    if a.record:
        with open(a.record, "w") as f:
            json.dump(record, f, indent=2, default=str)
            f.write("\n")


if __name__ == "__main__":
    main()
