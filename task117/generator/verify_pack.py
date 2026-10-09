"""Independent verifier for the task117 pack.

    python3 task117/generator/verify_pack.py <target dir> [--meta <metadata.json>]

Reads only the shipped files. Shares no code with the generator: it parses the rules out of the
rate schedule, the forecasting standard, the planning guide and the service agreement, rebuilds
every session's car through the permit and vehicle checks, replays the sessions on its own
quarter-hour arithmetic, and recomputes every rung, every rival cell, the calibration family,
the twin pair and both audits. Exit status 0 only when every figure matches."""
from __future__ import annotations

import itertools
import json
import math
import os
import re
import sys
from datetime import date, datetime, timedelta, timezone
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd
import pyarrow.parquet as pq

LA = ZoneInfo("America/Los_Angeles")
Q = 900

EXPECTED = {
    "answer": 109.088, "filed": 110, "split": (56.448, 52.64), "binding": "2026-01-21T12:00:00-08:00",
    "monthly": [109.088, 75.936, 81.760, 74.144, 76.832, 87.136, 80.192, 66.752, 95.200, 86.016, 79.072, 96.880],
    "rung0": 412.16, "rung1": 309.12, "rung2": 90.16, "rung3": 86.24, "own_date": 148.512,
    "own_date_binding": "2026-02-17T12:00:00-08:00", "planners": 225,
    "held": 129.136, "held_binding": "2026-12-08T12:00:00-08:00", "held_split": (82.88, 46.256),
    "rung4_records": 160.7424, "rung4_records_binding": "2026-01-21T12:00:00-08:00",
    "rung5_records": 155.35744, "rung5_records_binding": "2026-12-08T12:00:00-08:00",
    "b3_forecast": [112, 122, 105, 102, 102, 83, 83, 82, 107, 114, 121, 114],
    "b3_miss": [-0.9, 1.7, -1.9, 1.0, -1.0, 2.5, -1.2, 1.2, -1.8, 1.8, 1.7, -0.9],
    "b1": {"CP-N": [1069, 908, 910, 725, 594, 592, 599, 691, 791, 914, 1063, 1133],
           "CP-S": [934, 799, 794, 633, 520, 517, 523, 606, 693, 799, 932, 990]},
}
AS_OF = date(2027, 1, 25)
FAILS = []
OUT = {}


def check(name, cond, detail=""):
    print(("ok   " if cond else "FAIL ") + name + ("" if cond else f"  [{detail}]"))
    if not cond:
        FAILS.append(name)


def near(a, b, tol=1e-6):
    return abs(a - b) <= tol


# ------------------------------------------------------------------------------ reading the rules
def pdf_text(path):
    from pypdf import PdfReader
    return "\n".join(p.extract_text() for p in PdfReader(path).pages)


def rules(tdir):
    rs = pdf_text(os.path.join(tdir, "nspl_schedule_26_ev_charging_service.pdf"))
    flat = re.sub(r"\s+", " ", rs)
    m = re.search(r"fifteen-minute interval beginning at 12:00 noon through 7:45 p\.m\., Monday through Friday", flat)
    assert m, "billing window clause"
    toks = [t.strip() for t in rs.split("Observed dates:")[1].split("\n") if t.strip()]
    hol = set()
    names = ["New Year's Day", "Memorial Day", "Independence Day", "Labor Day", "Thanksgiving Day", "Christmas Day"]
    years = [2024, 2025, 2026, 2027, 2028]
    for name in names:
        i = toks.index(name)
        for y, cell in zip(years, toks[i + 1:i + 6]):
            mm = re.match(r"([A-Z][a-z]{2}) (\d{1,2})(?:, (\d{4}))?$", cell)
            assert mm, cell
            yy = int(mm.group(3)) if mm.group(3) else y
            hol.add(datetime.strptime(f"{mm.group(1)} {mm.group(2)} {yy}", "%b %d %Y").date())
    st = pdf_text(os.path.join(tdir, "fes-07_load_forecasting_standard_rev4.pdf"))
    fac = []
    for mm in re.finditer(r"(20\d\d)\n(1\.\d\d)\n([A-Z][a-z]+ \d{1,2}, \d{4})", st):
        fac.append((float(mm.group(2)), datetime.strptime(mm.group(3), "%B %d, %Y").date()))
    assert len(fac) >= 5, fac
    guide = pdf_text(os.path.join(tdir, "nspl_new_service_planning_guide_2026_sec7.pdf"))
    div = {}
    for mm in re.finditer(r"(\d+) to (\d+)\n(\d\.\d\d)", guide):
        div[(int(mm.group(1)), int(mm.group(2)))] = float(mm.group(3))
    from docx import Document
    d = Document(os.path.join(tdir, "civic_center_ev_service_agreement_draft.docx"))
    cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
    new_kw = {float(x.split()[0]) for x in cells if re.match(r"^\d+\.\d kW$", x)}
    units = sum(int(x) for x in cells if re.fullmatch(r"\d+", x))
    assert len(new_kw) == 1
    return {"holidays": hol, "factors": sorted(fac, key=lambda x: x[1]), "diversity": div,
            "new_kw": new_kw.pop(), "units": units}


def factor_on(factors, d):
    return [f for f, a in factors if a <= d][-1]


# ------------------------------------------------------------------------------ loading the pack
READ_LIST = ["settled_sessions_2024-2026.csv", "restatement_decisions_2025.csv", "station_register.csv",
             "session_intervals_2024-2026.parquet", "gateway_b_sessions_jan-apr2024.csv", "ev_permit_registry.csv",
             "permit_vehicle_checks.csv", "vehicle_reference_list.csv", "city_fleet_roster.csv",
             "fleet_card_ev_transactions_2024-2026.csv", "curbline_courtesy_sessions_civic_decks_2024-2026.csv",
             "ev_permit_charging_statements_2026.csv",
             "deck_panel_circuit_schedule.csv", "deck_submeter_nameplates.csv", "deck_panel_meter_log_2024-2026.xlsx",
             "nspl_schedule_26_ev_charging_service.pdf", "fes-07_load_forecasting_standard_rev4.pdf",
             "nspl_new_service_planning_guide_2026_sec7.pdf", "civic_center_ev_service_agreement_draft.docx"]
EPOCH = pd.Timestamp("1970-01-01", tz="UTC")


def epoch_s(ser):
    """Seconds since the epoch, whatever resolution pandas inferred."""
    return ((pd.to_datetime(ser, utc=True) - EPOCH) // pd.Timedelta(seconds=1)).astype("int64")


def load(tdir):
    j = lambda f: os.path.join(tdir, f)  # noqa: E731
    hdr = pd.read_csv(j("settled_sessions_2024-2026.csv"), dtype={"permit_no": str, "fleet_card": str})
    hdr["permit_no"] = hdr["permit_no"].fillna("")
    hdr["fleet_card"] = hdr["fleet_card"].fillna("")
    hdr["t0"] = epoch_s(hdr["plug_in"])
    hdr["t1"] = epoch_s(hdr["plug_out"])
    hdr["d"] = [datetime.fromisoformat(x).date() for x in hdr["plug_in"]]
    dec = pd.read_csv(j("restatement_decisions_2025.csv"))
    reg = pd.read_csv(j("station_register.csv"), parse_dates=["in_service_from", "in_service_to"])
    sp = pq.read_table(j("session_intervals_2024-2026.parquet")).to_pandas()
    sp["q"] = epoch_s(sp["interval_start"])
    gw = pd.read_csv(j("gateway_b_sessions_jan-apr2024.csv"))
    fc = pd.read_csv(j("fleet_card_ev_transactions_2024-2026.csv"))
    ct = pd.read_csv(j("curbline_courtesy_sessions_civic_decks_2024-2026.csv"), dtype={"Permit": str})
    ct["t0"] = [int(datetime.strptime(x, "%Y-%m-%d %H:%M:%S").replace(tzinfo=LA).timestamp()) for x in ct["Connected"]]
    ct["t1"] = [int(datetime.strptime(x, "%Y-%m-%d %H:%M:%S").replace(tzinfo=LA).timestamp()) for x in ct["Disconnected"]]
    ct["d"] = [datetime.strptime(x, "%Y-%m-%d %H:%M:%S").date() for x in ct["Connected"]]
    ct["station_id"] = ct["Station"]
    stm = pd.read_csv(j("ev_permit_charging_statements_2026.csv"))
    permits = pd.read_csv(j("ev_permit_registry.csv"))
    checks = pd.read_csv(j("permit_vehicle_checks.csv"), dtype={"trim": str})
    checks["trim"] = checks["trim"].fillna("")
    ref = pd.read_csv(j("vehicle_reference_list.csv"), dtype={"trim": str})
    ref["trim"] = ref["trim"].fillna("")
    fleet = pd.read_csv(j("city_fleet_roster.csv"), dtype={"trim": str})
    fleet["trim"] = fleet["trim"].fillna("")
    sched = pd.read_csv(j("deck_panel_circuit_schedule.csv"), parse_dates=["effective_from", "effective_to"])
    plates = pd.read_csv(j("deck_submeter_nameplates.csv"))
    from openpyxl import load_workbook
    wb = load_workbook(j("deck_panel_meter_log_2024-2026.xlsx"), read_only=True)

    def sheet(ws):
        rows = list(ws.iter_rows(values_only=True))
        hi = [i for i, r in enumerate(rows) if r and r[0] == "Read date"][0]
        df = pd.DataFrame([r for r in rows[hi + 1:] if r and r[0] is not None], columns=list(rows[hi]))
        df["Read date"] = pd.to_datetime(df["Read date"]).dt.date
        return df
    log = sheet(wb.worksheets[0])
    fixes = sheet(wb["Corrections"])
    return dict(hdr=hdr, dec=dec, reg=reg, sp=sp, gw=gw, fc=fc, ct=ct, stm=stm, permits=permits, checks=checks, ref=ref,
                fleet=fleet, sched=sched, plates=plates, log=log, fixes=fixes,
                log_path=j("deck_panel_meter_log_2024-2026.xlsx"))


def dated_join(P, frame, tcol="d", with_rating=False):
    """Garage and position of each row by the register's dated assignment."""
    reg = P["reg"].copy()
    reg["to"] = reg["in_service_to"].fillna(pd.Timestamp("2099-12-31"))
    m = frame[["station_id", tcol]].reset_index().merge(reg, on="station_id", how="left")
    ok = (m["in_service_from"].dt.date <= m[tcol]) & (m["to"].dt.date >= m[tcol])
    m = m[ok].set_index("index")
    assert not m.index.duplicated().any()
    if with_rating:
        return m["garage"].reindex(frame.index), m["position"].reindex(frame.index), m["rating_kw"].reindex(frame.index)
    return m["garage"].reindex(frame.index), m["position"].reindex(frame.index)


def of_record(P):
    """Session rows of record: accepted versions, one row per authorization code."""
    h = P["hdr"]
    dec = P["dec"]
    acc = dec[dec["decision"] == "ACCEPTED"].groupby("session_id")["version"].max()
    rest = set(dec["session_id"])
    keep_ver = h["session_id"].map(acc).fillna(1).astype(int)
    restated = h["session_id"].isin(rest)
    sel = (~restated & (h["version"] == 1)) | (restated & (h["version"] == keep_ver))
    r = h[sel]
    r = r.sort_values(["session_id"]).drop_duplicates("auth_code", keep="first")
    return r


def charges(fr):
    """One row per charge: records at one station on one permit or card that meet end to start are one charge."""
    fr = fr.sort_values(["station_id", "t0"], kind="mergesort").copy()
    who = fr["permit_no"].where(fr["permit_no"] != "", fr["fleet_card"])
    cont = (fr["t0"] == fr["t1"].shift()) & (fr["station_id"] == fr["station_id"].shift()) & (who == who.shift())
    cid = (~cont).cumsum().to_numpy()
    g = fr.groupby(cid, sort=False)
    first = g.head(1).copy()
    first["t0"] = g["t0"].min().to_numpy()
    first["t1"] = g["t1"].max().to_numpy()
    first["kwh_delivered"] = np.round(g["kwh_delivered"].sum().to_numpy(), 3)
    first["n_rec"] = g.size().to_numpy()
    return first


def of_record_asof(P, when):
    """Session rows as they stood on a date: a restated version counts once it was decided ACCEPTED by then."""
    h = P["hdr"]
    dec = P["dec"].copy()
    dec["decd"] = pd.to_datetime(dec["decided_on"]).dt.date
    acc = dec[(dec["decision"] == "ACCEPTED") & (dec["decd"] <= when)].groupby("session_id")["version"].max()
    keep_ver = h["session_id"].map(acc).fillna(1).astype(int)
    r = h[h["version"] == keep_ver]
    return r.sort_values(["session_id"]).drop_duplicates("auth_code", keep="first")


POOL = "Larch County Fleet Services"


def handoff_prev(ch, pool, window=600, need_pool=True):
    """For each charge, the charge it follows on a hand-off: the next charge at a unit is a pool car going on within
    the window of a pool car coming off."""
    o = ch.sort_values(["station_id", "t0"], kind="mergesort")
    st, t0, t1 = o["station_id"].to_numpy(), o["t0"].to_numpy(np.int64), o["t1"].to_numpy(np.int64)
    pm = o["permit_no"].isin(pool).to_numpy() if need_pool else np.ones(len(o), bool)
    gap = t0[1:] - t1[:-1]
    h = (st[1:] == st[:-1]) & pm[1:] & pm[:-1] & (gap > 0) & (gap <= window)
    prev = {}
    idx = o.index.to_numpy()
    for k in np.flatnonzero(h):
        prev[idx[k + 1]] = idx[k]
    return prev


def forward_starts(ch, rate, prev, old_kw, lag="own"):
    """Start of each charge on the new units: a hand-off follower goes on once the car ahead of it has finished at its
    new draw, after the wait it had behind that car's finish on the old units (or the mean wait, or none)."""
    start = ch["t0"].astype(float)
    fin_old = start + ch["kwh_delivered"].astype(float) / old_kw * 3600.0
    waits = {i: start[i] - fin_old[p] for i, p in prev.items()}
    mean = float(np.mean(list(waits.values()))) if waits else 0.0
    rate = pd.Series(rate, index=ch.index).astype(float)
    new = {}
    for i in ch.sort_values("t0", kind="mergesort").index:
        if i not in prev:
            new[i] = start[i]
            continue
        p = prev[i]
        w = waits[i] if lag == "own" else mean if lag == "mean" else 0.0
        new[i] = new[p] + ch.at[p, "kwh_delivered"] / rate[p] * 3600.0 + w
    return pd.Series(new).reindex(ch.index)



# ------------------------------------------------------------------------------ cars
_RATING = {}


def car_ratings(P, s, as_of=None):
    ref = P["ref"]
    ch = P["checks"].copy()
    ch["checked_on"] = pd.to_datetime(ch["checked_on"]).dt.date

    def rating(mk, md, tr, yr):
        key = (mk.upper(), md.upper(), tr.upper(), yr)
        if key not in _RATING:
            hit = ref[(ref["make"].str.upper() == key[0]) & (ref["model"].str.upper() == key[1])
                      & (ref["model_year_from"] <= yr) & (ref["model_year_to"] >= yr)
                      & ((ref["trim"] == "") | (ref["trim"].str.upper() == key[2]))]
            assert len(hit) == 1, key
            _RATING[key] = float(hit["onboard_charger_kw"].iloc[0])
        return _RATING[key]
    by_permit = {k: g.sort_values("checked_on") for k, g in ch.groupby("permit_no")}
    fl = P["fleet"].set_index("fleet_card")
    out = []
    for pn, fc, d in zip(s["permit_no"], s["fleet_card"], s["d"]):
        if pn:
            g = by_permit[pn]
            g = g[g["checked_on"] <= (as_of or d)]
            c = g.iloc[-1]
            out.append(rating(c["make"], c["model"], c["trim"], int(c["model_year"])))
        else:
            f = fl.loc[fc]
            out.append(rating(f["make"], f["model"], f["trim"], int(f["model_year"])))
    return pd.Series(out, index=s.index)


# ------------------------------------------------------------------------------ quarter-hour arithmetic
class Grid:
    def __init__(self, holidays):
        self.t0 = int(datetime(2023, 12, 25, tzinfo=LA).timestamp())
        self.t1 = int(datetime(2027, 1, 4, tzinfo=LA).timestamp())
        self.n = (self.t1 - self.t0) // Q
        loc = pd.to_datetime(self.t0 + Q * np.arange(self.n), unit="s", utc=True).tz_convert(LA)
        self.year, self.month = loc.year.to_numpy(), loc.month.to_numpy()
        dates = loc.date
        hol = np.array([x in holidays for x in dates])
        mins = loc.hour.to_numpy() * 60 + loc.minute.to_numpy()
        self.bill = (loc.weekday.to_numpy() < 5) & (mins >= 12 * 60) & (mins <= 19 * 60 + 45) & ~hol
        self.loc = loc

    def blocks(self, start, energy, rate):
        """Average kW per quarter-hour of constant-rate blocks, built quarter by quarter."""
        out = np.zeros(self.n)
        start = np.asarray(start, float)
        end = start + np.asarray(energy, float) / np.asarray(rate, float) * 3600.0
        rate = np.asarray(rate, float) * np.ones_like(start)
        k0 = np.floor((start - self.t0) / Q).astype(int)
        k1 = np.floor((end - self.t0 - 1e-9) / Q).astype(int)
        cnt = k1 - k0 + 1
        rep = np.repeat(np.arange(len(start)), cnt)
        k = k0[rep] + (np.arange(cnt.sum()) - np.repeat(np.cumsum(cnt) - cnt, cnt))
        qs = self.t0 + Q * k
        ov = np.minimum(qs + Q, end[rep]) - np.maximum(qs, start[rep])
        np.add.at(out, k, rate[rep] * ov / Q)
        return out

    def readings_load(self, q, kwh):
        out = np.zeros(self.n)
        np.add.at(out, (np.asarray(q) - self.t0) // Q, 4 * np.asarray(kwh))
        return out

    def peak(self, load, year=2026, month=None, mask=None):
        m = (self.bill if mask is None else mask) & (self.year == year)
        if month:
            m &= self.month == month
        idx = np.flatnonzero(m)
        j = idx[np.argmax(load[idx])]
        return float(load[j]), self.t0 + Q * j


def iso(t):
    return datetime.fromtimestamp(int(t), LA).isoformat()


# ------------------------------------------------------------------------------ the verification
def main():
    tdir = sys.argv[1]
    meta = None
    if "--meta" in sys.argv:
        meta = json.load(open(sys.argv[sys.argv.index("--meta") + 1]))
    R = rules(tdir)
    P = load(tdir)
    G = Grid(R["holidays"])
    growth = factor_on(R["factors"], date(2027, 1, 25))
    check("V01 the factor in force for a forecast made in January 2027 is 1.12", growth == 1.12, growth)

    rec = of_record(P)
    rec = rec.copy()
    rec["garage"], rec["position"] = dated_join(P, rec)
    decks = ["Civic Center North Deck", "Civic Center South Deck"]
    recs = rec[rec["garage"].isin(decks) & (pd.to_datetime(rec["plug_in"].str[:10]).dt.year == 2026)].copy()
    pop = charges(recs)
    check("V01a the 2026 deck records join into charges: every record meeting another end to start at its station on "
          "the same permit or card is one charge, and no record overlaps another at a station",
          int((pop["n_rec"] > 1).sum()) > 2000 and int(pop["n_rec"].max()) == 2, (len(recs), len(pop)))
    pop["car"] = car_ratings(P, pop)
    pop["car27"] = car_ratings(P, pop, as_of=AS_OF)
    recs["car"] = car_ratings(P, recs)
    recs["car27"] = car_ratings(P, recs, as_of=AS_OF)
    old_kw = float(P["reg"].loc[P["reg"]["garage"].isin(decks), "rating_kw"].unique()[0])
    check("V02 every deck unit is rated 6.6 kW today and the new units 11.5 kW", old_kw == 6.6 and R["new_kw"] == 11.5)
    per = P["permits"].set_index("permit_no")
    pool = set(per.index[per["holder"] == POOL])
    prev = handoff_prev(pop, pool)

    def replay(rate, frame=None, starts=None):
        frame = pop if frame is None else frame
        rate = pd.Series(rate, index=frame.index)
        t0 = frame["t0"] if starts is None else starts
        return {g: G.blocks(t0[frame["garage"] == g], frame.loc[frame["garage"] == g, "kwh_delivered"],
                            rate[frame["garage"] == g]) for g in decks}

    def replay_fwd(rate, frame=None, prv=None, lag="own"):
        frame = pop if frame is None else frame
        prv = prev if prv is None else prv
        return replay(rate, frame, forward_starts(frame, rate, prv, old_kw, lag))
    # the answer: every 2026 charge replayed at the smaller of the new rating and the onboard rating of the car its
    # permit carries into the contract year, the county pool's hand-offs going on once the car ahead has finished
    r27 = np.minimum(R["new_kw"], pop["car27"])
    percar = replay_fwd(r27)
    tot = percar[decks[0]] + percar[decks[1]]
    v, t = G.peak(tot)
    ans = growth * v
    OUT["answer"] = ans
    check("V03 the answer recomputes: 109.088 kW at 12:00 on 21 January 2026",
          near(ans, EXPECTED["answer"], 1e-6) and iso(t) == EXPECTED["binding"], (ans, iso(t)))
    filed = int(math.floor(ans / 5 + 0.5) * 5)
    check("V04 files 110 kW to the nearest 5 kW (and rounded up)", filed == 110 and math.ceil(ans / 5) * 5 == 110)
    j = (t - G.t0) // Q
    split = (growth * percar[decks[0]][j], growth * percar[decks[1]][j])
    OUT["split"] = split
    check("V05 deck split 56.4 / 52.6 kW", near(split[0], 56.448, 1e-6) and near(split[1], 52.64, 1e-6), split)
    mon = [growth * G.peak(tot, month=m)[0] for m in range(1, 13)]
    OUT["monthly"] = mon
    check("V06 the twelve contract months recompute", all(near(a, b, 1e-6) for a, b in zip(mon, EXPECTED["monthly"])), mon)
    check("V07 the month that sets the figure is January (contract month January 2028)",
          max(range(12), key=lambda i: mon[i]) == 0)
    # the hand-offs, read from the export and the permit registry alone
    allc = charges(rec[rec["garage"].isin(decks)].copy())
    prev_all = handoff_prev(allc, pool)
    o = allc.sort_values(["station_id", "t0"], kind="mergesort")
    st_o, t0_o, t1_o = o["station_id"].to_numpy(), o["t0"].to_numpy(np.int64), o["t1"].to_numpy(np.int64)
    gap_o = t0_o[1:] - t1_o[:-1]
    is_h = np.array([i in prev_all for i in o.index.to_numpy()[1:]])
    others = (st_o[1:] == st_o[:-1]) & ~is_h
    fol = allc.loc[list(prev_all)]
    head = allc.loc[[prev_all[i] for i in prev_all]]
    idle = (head["t1"].to_numpy() - (head["t0"].to_numpy() + head["kwh_delivered"].to_numpy() / old_kw * 3600)) / 60
    lagm = (fol["t0"].to_numpy() - head["t1"].to_numpy()) / 60
    OUT["handoffs"] = (len(prev_all), len(prev))
    check("V07k the county pool's hand-offs: 282 in the export from 2024 to 2026, 101 in 2026, every one at a North Deck "
          "unit; the car ahead came off 6 to 20 minutes after its last charging second and the next went on 1 to 5.4 "
          "minutes later; every other pair of consecutive charges at a deck unit sits 42 minutes or more apart",
          len(prev_all) == 282 and len(prev) == 101 and (fol["garage"] == decks[0]).all()
          and idle.min() >= 6.0 and idle.max() <= 20.0 and lagm.min() >= 1.0 and lagm.max() <= 5.4
          and gap_o[others].min() >= 42 * 60,
          (len(prev_all), len(prev), round(float(idle.min()), 2), round(float(idle.max()), 2), float(gap_o[others].min()) / 60))
    same = all(handoff_prev(pop, pool, w_, need) == prev for w_ in (360, 2520) for need in (True, False))
    conv = {}
    for lag in ("mean", "zero"):
        lo = replay_fwd(r27, lag=lag)
        tl = lo[decks[0]] + lo[decks[1]]
        conv[lag] = (growth * G.peak(tl)[0], G.peak(tl)[1], [growth * G.peak(tl, month=m)[0] for m in range(1, 13)])
    check("V07l the hand-offs converge: every window from 6 to 42 minutes, with or without the pool condition, selects "
          "the same hand-offs; the mean wait and no wait give the answer, its quarter-hour and every month",
          same and all(near(c[0], ans, 1e-9) and c[1] == t and all(near(a, b, 1e-9) for a, b in zip(c[2], mon))
                       for c in conv.values()), {k: c[0] for k, c in conv.items()})
    held = replay(r27)
    th = held[decks[0]] + held[decks[1]]
    vh, tth = G.peak(th)
    jh = (tth - G.t0) // Q
    OUT["held"] = growth * vh
    hsplit = (growth * held[decks[0]][jh], growth * held[decks[1]][jh])
    jd = int((int(datetime(2026, 12, 8, 12, tzinfo=LA).timestamp()) - G.t0) // Q)
    check("V07m every charge replayed from its 2026 start (the hand-offs held) lands 129.136 kW at 12:00 on 8 December "
          "2026 (82.9 / 46.3), filed 130; re-timing the hand-offs takes December's 12:00 quarter-hour from 115.3 to 86.5 kW",
          near(growth * vh, EXPECTED["held"], 1e-6) and iso(tth) == EXPECTED["held_binding"]
          and near(hsplit[0], 82.88, 1e-6) and near(hsplit[1], 46.256, 1e-6)
          and near(th[jd], 115.3, 1e-6) and near(tot[jd], 86.5, 1e-6), (growth * vh, iso(tth), th[jd], tot[jd]))
    closed_fwd = replay_fwd(np.full(len(pop), old_kw))
    closed_held = replay(np.full(len(pop), old_kw))
    check("V07n the closed record is blind to the hand-offs: at 6.6 kW every re-timed charge starts when it started",
          all(np.allclose(closed_fwd[g], closed_held[g], atol=1e-9) for g in decks))
    # the renewal, read from the checks alone
    ch = P["checks"].copy()
    ch["checked_on"] = pd.to_datetime(ch["checked_on"]).dt.date
    jan = ch[ch["checked_on"] >= date(2027, 1, 1)].set_index("permit_no")
    before = ch[ch["checked_on"] < date(2027, 1, 1)].sort_values("checked_on").groupby("permit_no").last()
    changed = sorted(p for p in jan.index if jan.loc[p, "vin"] != before.loc[p, "vin"])
    check("V07a the January 2027 renewal: eighteen county permits moved from 2020 to 2023 Bolt EVs, checked by 15 January",
          len(changed) == 18 and set(per.loc[changed, "holder"]) == {POOL}
          and set(before.loc[changed, "model_year"]) == {2020} and set(jan.loc[changed, "model_year"]) == {2023}
          and jan["checked_on"].max() <= date(2027, 1, 15), (len(changed), jan["checked_on"].max()))
    diff = pop[pop["car"] != pop["car27"]]
    check("V07b the two vehicle joins differ only on the renewed permits' sessions, 7.2 kW then and 11.0 kW in the "
          "contract year", set(diff["permit_no"]) <= set(changed) and (diff["car"] == 7.2).all()
          and (diff["car27"] == 11.0).all() and len(diff) == int(pop["permit_no"].isin(changed).sum()), len(diff))
    # each session's car on its own date
    r26c = np.minimum(R["new_kw"], pop["car"])
    own_date = replay_fwd(r26c)
    own_held = replay(r26c)
    v4, t4 = G.peak(own_date[decks[0]] + own_date[decks[1]])
    v4h, t4h = G.peak(own_held[decks[0]] + own_held[decks[1]])
    OUT["own_date"] = growth * v4
    check("V07c each session's own car lands 148.512 kW at 12:00 on 17 February 2026, the hand-offs re-timed or held, "
          "filed 150", near(growth * v4, EXPECTED["own_date"], 1e-6) and iso(t4) == EXPECTED["own_date_binding"]
          and near(v4h, v4, 1e-9) and t4h == t4 and int(math.floor(growth * v4 / 5 + 0.5) * 5) == 150,
          (growth * v4, iso(t4)))
    # the natural path replays settlement records: rung 4 (own-date cars) and rung 5 (contract-year cars)
    prev_rec = handoff_prev(recs, pool)
    r5 = replay(np.minimum(R["new_kw"], recs["car27"]), recs)
    v5, t5 = G.peak(r5[decks[0]] + r5[decks[1]])
    r4 = replay(np.minimum(R["new_kw"], recs["car"]), recs)
    v4r, t4r = G.peak(r4[decks[0]] + r4[decks[1]])
    r5f = replay_fwd(np.minimum(R["new_kw"], recs["car27"]), recs, prev_rec)
    r4f = replay_fwd(np.minimum(R["new_kw"], recs["car"]), recs, prev_rec)
    OUT["rung5_records"], OUT["rung4_records"] = growth * v5, growth * v4r
    check("V07d record by record, rung 5 lands 155.357 kW at 12:00 on 8 December 2026 and rung 4 160.742 kW at 12:00 on "
          "21 January 2026, and re-timing the hand-offs on records moves neither",
          near(growth * v5, EXPECTED["rung5_records"], 1e-5) and iso(t5) == EXPECTED["rung5_records_binding"]
          and near(growth * v4r, EXPECTED["rung4_records"], 1e-5) and iso(t4r) == EXPECTED["rung4_records_binding"]
          and near(G.peak(r5f[decks[0]] + r5f[decks[1]])[0], v5, 1e-9)
          and near(G.peak(r4f[decks[0]] + r4f[decks[1]])[0], v4r, 1e-9), (growth * v5, iso(t5), growth * v4r, iso(t4r)))
    j5, j4 = (t5 - G.t0) // Q, (t4r - G.t0) // Q
    sp5 = (growth * r5[decks[0]][j5], growth * r5[decks[1]][j5])
    sp4 = (growth * r4[decks[0]][j4], growth * r4[decks[1]][j4])
    check("V07h record by record: rung 5 82.88 and 72.477 kW at its quarter-hour, rung 4 88.704 and 72.038 kW",
          near(sp5[0], 82.88, 1e-6) and near(sp5[1], 72.47744, 1e-5) and near(sp4[0], 88.704, 1e-6)
          and near(sp4[1], 72.0384, 1e-5), (sp5, sp4))
    mon5 = [growth * G.peak(r5[decks[0]] + r5[decks[1]], month=m)[0] for m in range(1, 13)]
    check("V07e rung 5 differs from the answer at whole kW in every month", all(round(a) != round(b) for a, b in
                                                                              zip(mon5, mon)), mon5)
    # the statements bill per charge: only the joined charges reproduce every permit-month
    stm = P["stm"].set_index(["permit_no", "period"])
    pr = recs[recs["permit_no"] != ""].copy()
    pr["period"] = pr["plug_in"].str[:7]
    pc = charges(pr)
    n_rec = pr.groupby(["permit_no", "period"]).size()
    n_chg = pc.groupby(["permit_no", "period"]).size()
    n_over = pr.groupby(["permit_no", "period"]).apply(lambda g: g.groupby(["station_id", "d"]).ngroups)
    kwh = pr.groupby(["permit_no", "period"])["kwh_delivered"].sum().round(3)
    idx = stm.index
    check("V07f the permit statements reproduce on every permit-month from the joined charges (counts and kWh); "
          "records as charges miss most months, merging every same-day record at a unit misses some",
          (n_chg.reindex(idx).fillna(0).astype(int) == stm["charges"]).all()
          and (np.abs(kwh.reindex(idx).fillna(0) - stm["kwh"]) < 0.0005).all()
          and int((n_rec.reindex(idx).fillna(0).astype(int) != stm["charges"]).sum()) >= 0.5 * len(idx)
          and int((n_over.reindex(idx).fillna(0).astype(int) != stm["charges"]).sum()) >= 10, len(idx))
    allr = rec.copy()
    allr["who"] = allr["permit_no"].where(allr["permit_no"] != "", allr["fleet_card"])
    allr = allr[allr["who"] != ""].sort_values(["station_id", "who", "t0"], kind="mergesort")
    same_k = ((allr["station_id"].to_numpy()[1:] == allr["station_id"].to_numpy()[:-1])
              & (allr["who"].to_numpy()[1:] == allr["who"].to_numpy()[:-1]))
    gaps_k = (allr["t0"].to_numpy()[1:] - allr["t1"].to_numpy()[:-1])[same_k]
    check("V07i at one unit under one permit or card, records either meet end to start or sit 42 minutes or more apart, "
          "and never overlap, 2024 to 2026", int((gaps_k < 0).sum()) == 0 and gaps_k[gaps_k > 0].min() >= 42 * 60,
          (int((gaps_k == 0).sum()), float(gaps_k[gaps_k > 0].min()) / 60))
    off = pd.to_datetime(rec["plug_in"].str[:10])
    off_day = (off.dt.weekday >= 5) | off.dt.date.isin(set(R["holidays"]))
    by_g = off_day.groupby(rec["garage"]).sum()
    check("V07j the settlement export carries no weekend or holiday session at either deck, and some at each of the other "
          "six garages", int(by_g[decks].sum()) == 0 and int((by_g.drop(decks) > 0).sum()) == 6, by_g.to_dict())
    # merging every same-day record of a permit or card at a unit also joins the genuine replugs; on the call it
    # converges, hand-offs included
    ov = recs.copy()
    ov["who"] = ov["permit_no"].where(ov["permit_no"] != "", ov["fleet_card"])
    ov["day"] = ov["plug_in"].str[:10]
    om = ov.sort_values("t0", kind="mergesort").groupby(["station_id", "who", "day"], as_index=False, sort=False).agg(
        t0=("t0", "min"), t1=("t1", "max"), kwh_delivered=("kwh_delivered", "sum"), garage=("garage", "first"),
        car27=("car27", "first"), permit_no=("permit_no", "first"))
    lo = replay_fwd(np.minimum(R["new_kw"], om["car27"]), om, handoff_prev(om, pool))
    tot_om = lo[decks[0]] + lo[decks[1]]
    v_om, t_om = G.peak(tot_om)
    mon_om = [growth * G.peak(tot_om, month=m)[0] for m in range(1, 13)]
    check("V07g merging every same-day record of a permit or card at a unit leaves the answer, its quarter-hour and "
          "every month unchanged", len(om) < len(pop) and near(growth * v_om, ans, 1e-9) and t_om == t
          and all(near(a, b, 1e-9) for a, b in zip(mon_om, mon)), (len(om), len(pop), growth * v_om))
    # rungs
    log = P["log"]
    l26 = log[pd.to_datetime(log["Read date"]).dt.year == 2026]
    per_m = l26.groupby(["Read date", "Panel"])["Max demand (kW)"].max().unstack().sum(axis=1)
    r0 = float(per_m.max()) * growth * R["new_kw"] / old_kw
    check("V08 rung 0 recomputes from the panel log: 412.16", near(r0, EXPECTED["rung0"], 1e-6), r0)
    sp = P["sp"]
    deck_ids = set(recs["session_id"])
    closed_rows = sp[sp["session_id"].isin(deck_ids)]
    vkey = recs.set_index("session_id")["version"]
    closed_rows = closed_rows[closed_rows["version"] == closed_rows["session_id"].map(vkey)]
    closed = G.readings_load(closed_rows["q"], closed_rows["kwh"])
    c_peak, c_t = G.peak(closed)
    r1 = c_peak * growth * R["new_kw"] / old_kw
    check("V09 rung 1 recomputes from the quarter-hour readings: 158.4 closed, 309.12", near(c_peak, 158.4, 1e-6)
          and near(r1, EXPECTED["rung1"], 1e-6), (c_peak, r1))
    r115 = replay_fwd(np.full(len(pop), R["new_kw"]))
    v2, t2 = G.peak(r115[decks[0]] + r115[decks[1]])
    check("V10 rung 2 recomputes: 90.16 at 12:00 on 10 June 2026", near(growth * v2, EXPECTED["rung2"], 1e-6)
          and iso(t2) == "2026-06-10T12:00:00-07:00", (growth * v2, iso(t2)))
    lib = rec[(rec["position"] == "L-09")]
    obs = observed(P, lib)
    lib_rate = float(obs[lib["fleet_card"].map(P["fleet"].set_index("fleet_card")["department"]) ==
                         "Parking Enforcement"].median())
    r110 = replay_fwd(np.full(len(pop), lib_rate))
    v3, t3 = G.peak(r110[decks[0]] + r110[decks[1]])
    check("V11 rung 3 recomputes at the Library unit's van draw: 86.24", near(growth * v3, EXPECTED["rung3"], 1e-6),
          growth * v3)
    all_hours = float(closed[(G.year == 2026)].max())
    check("V11a the closed all-hours maximum is the registers' 211.2 kW, the billing-hours maximum 158.4",
          near(all_hours, 211.2, 1e-6) and c_peak < all_hours, (all_hours, c_peak))
    shift = np.roll(G.bill, -1)
    mon_shift = [growth * G.peak(tot, month=m, mask=shift)[0] for m in range(1, 13)]
    check("V11b reading the stamps as interval ends returns every monthly figure",
          all(near(a, b, 1e-6) for a, b in zip(mon, mon_shift)))
    own = {g: G.peak(percar[g])[1] for g in decks}
    check("V11c each deck's own maximum falls in the binding quarter-hour", all(x == t for x in own.values()))
    # grid: one violation of the answer's construction per cell
    cells = {"draw 6.6": growth * c_peak, "growth left off": v, "cars on their own dates": growth * v4,
             "the hand-offs held at their 2026 starts": growth * vh, "settlement records": growth * v5}
    veh = vehicles_at(P, AS_OF)
    favg = float(np.mean(np.minimum(R["new_kw"], veh["rating"])))
    davg = {g: float(np.mean(np.minimum(R["new_kw"], veh.loc[veh["deck"] == k, "rating"])))
            for g, k in zip(decks, ("North", "South"))}
    cells["fleet-average ratio, contract-year vehicles"] = growth * c_peak * favg / old_kw
    cells["fleet-average replay, contract-year vehicles"] = growth * G.peak(sum(replay_fwd(np.full(len(pop), favg)).values()))[0]
    cells["deck-average replay, contract-year vehicles"] = growth * G.peak(sum(replay_fwd(
        pop["garage"].map(davg).to_numpy()).values()))[0]
    cells["per car North, rating South"] = growth * G.peak(sum(replay_fwd(np.where(
        pop["garage"] == decks[0], r27, 11.5)).values()))[0]
    cells["per car South, rating North"] = growth * G.peak(sum(replay_fwd(np.where(
        pop["garage"] == decks[1], r27, 11.5)).values()))[0]
    units = R["units"]
    divf = [f for (a, b), f in R["diversity"].items() if a <= units <= b][0]
    planners = math.ceil(units * R["new_kw"] * divf / 5) * 5
    cells["planners' sizing"] = planners
    OUT["cells"] = cells
    check("V12 the planners' figure is 225 kW (32 x 11.5 x 0.60 to the next 5 kW above), 115 kW above ours",
          planners == 225 and planners - filed == 115, planners)
    check("V13 every single-error cell at least 10 per cent from the answer",
          all(abs(x / ans - 1) >= 0.10 for x in cells.values()), {k: round(x / ans - 1, 3) for k, x in cells.items()})
    veh26 = vehicles_at(P, date(2026, 12, 31))
    davg26 = {g: float(np.mean(np.minimum(R["new_kw"], veh26.loc[veh26["deck"] == k, "rating"])))
              for g, k in zip(decks, ("North", "South"))}
    favg26 = float(np.mean(np.minimum(R["new_kw"], veh26["rating"])))
    two = {"deck averages on 2026 vehicles": growth * G.peak(sum(replay_fwd(pop["garage"].map(davg26).to_numpy()).values()))[0],
           "fleet average on 2026 vehicles": growth * G.peak(sum(replay_fwd(np.full(len(pop), favg26)).values()))[0],
           "the hand-offs held, growth left off": vh, "own-date cars, growth left off": v4}
    check("V14 two-error cells: each deck's average on 2026 vehicles files 115, the fleet average on 2026 vehicles sits "
          "at least 40 per cent above, the hand-offs held with growth left off files 115, own-date cars with growth left "
          "off 135; none files 110", int(math.floor(two["deck averages on 2026 vehicles"] / 5 + 0.5) * 5) == 115
          and two["fleet average on 2026 vehicles"] >= 1.4 * ans
          and int(math.floor(two["the hand-offs held, growth left off"] / 5 + 0.5) * 5) == 115
          and int(math.floor(two["own-date cars, growth left off"] / 5 + 0.5) * 5) == 135
          and all(int(math.floor(x / 5 + 0.5) * 5) != 110 and math.ceil(x / 5) * 5 != 110 for x in two.values()),
          {k: round(x, 3) for k, x in two.items()})
    # calibration: the five-rule family on every session with an identified vehicle
    idd = rec[(rec["permit_no"] != "") | (rec["fleet_card"] != "")].copy()
    idd["obs"] = observed(P, idd)
    meas = idd.dropna(subset=["obs"]).copy()
    meas["car"] = car_ratings(P, meas)
    unit = dated_join(P, meas, with_rating=True)[2].astype(float)
    fam = {"composition": np.minimum(unit, meas["car"]), "rating": unit, "onboard": meas["car"],
           "derate": unit * 11.0 / 11.5, "fixed 11.0": unit.where(unit != 11.5, 11.0)}
    misses = {k: int((np.abs(x - meas["obs"]) > 0.01).sum()) for k, x in fam.items()}
    OUT["family"] = misses
    check("V15 the composition reproduces every measurable session; every rival misses a whole population",
          misses["composition"] == 0 and min(misses["rating"], misses["onboard"], misses["derate"], misses["fixed 11.0"]) >= 10,
          misses)
    pick = meas[(meas["position"] == "L-09") & (meas["car"] > 11.5)]
    check("V16 the Library unit: vans at 11.0 kW, the pickup at 11.5 kW in at least ten sessions",
          len(pick) >= 10 and (pick["obs"] == 11.5).all() and lib_rate == 11.0, (len(pick), lib_rate))
    dk = rec[rec["garage"].isin(decks)]
    minc = min(float(car_ratings(P, dk).min()), float(car_ratings(P, dk, as_of=AS_OF).min()))
    check("V17 every deck vehicle, then and in the contract year, accepts at least 7.2 kW, so no closed deck session "
          "shows its car's limit", minc >= 7.2, minc)
    model_err = session_model(P, rec)
    check("V17a the session model regenerates every closed reading (constant draw until delivered, then zero) to 0.001 kWh",
          model_err <= 0.0011, model_err)
    # twins
    loc = pd.to_datetime(pop["plug_in"].str[11:19], format="%H:%M:%S")
    arr = loc.dt.hour + loc.dt.minute / 60
    dw = (pop["t1"] - pop["t0"]) / 3600
    ks = {n: ks2(x[pop["garage"] == decks[0]], x[pop["garage"] == decks[1]])
          for n, x in (("arrival", arr), ("dwell", dw), ("energy", pop["kwh_delivered"]))}
    cnt = pop.groupby("garage").size()
    north_ids = recs.loc[recs["garage"] == decks[0], "session_id"]
    cl_n = G.readings_load(closed_rows.loc[closed_rows["session_id"].isin(north_ids), "q"],
                           closed_rows.loc[closed_rows["session_id"].isin(north_ids), "kwh"])
    check("V18 twin decks matched on counts and distributions; equal closed load at the answer's quarter-hour (72.6 kW "
          "each) and at 12:00 on 8 December 2026 (66.0 kW each)",
          abs(cnt.iloc[0] / cnt.iloc[1] - 1) <= 0.04 and max(ks.values()) < 0.07
          and near(cl_n[j], 72.6, 1e-6) and near(closed[j] - cl_n[j], 72.6, 1e-6)
          and near(cl_n[jh], 66.0, 1e-6) and near(closed[jh] - cl_n[jh], 66.0, 1e-6), (cnt.to_dict(), ks, cl_n[j]))
    # B3
    OUT["b3"] = b3(P, R, G, rec)
    # B1
    OUT["b1"] = b1(P, R, G, rec)
    if meta:
        check("V31 the two distractors are outside this verifier's read list, so no figure depends on them",
              not set(meta["distractor_files"]) & set(READ_LIST))
        files = sorted(os.listdir(tdir))
        check("V30 input gates: 10+ files, 3+ formats, 25,000+ rows, two distractors present and unread",
              len(files) >= 10 and len({os.path.splitext(f)[1] for f in files}) >= 3
              and pq.ParquetFile(os.path.join(tdir, "session_intervals_2024-2026.parquet")).metadata.num_rows >= 25000
              and len(meta["distractor_files"]) >= 2 and all(d in files for d in meta["distractor_files"]))
    print(json.dumps({k: v for k, v in OUT.items() if k in ("answer", "split", "held", "handoffs")}, default=str))
    print(f"\n{len(FAILS)} failure(s)")
    sys.exit(1 if FAILS else 0)


def ks2(a, b):
    a, b = np.sort(np.asarray(a, float)), np.sort(np.asarray(b, float))
    v = np.concatenate([a, b])
    return float(np.max(np.abs(np.searchsorted(a, v, "right") / len(a) - np.searchsorted(b, v, "right") / len(b))))


def observed(P, rows):
    """kW from full charging quarter-hours: the largest reading of the session times four."""
    sp = P["sp"]
    r = sp[sp["session_id"].isin(rows["session_id"])]
    key = rows.set_index("session_id")["version"]
    r = r[r["version"] == r["session_id"].map(key)]
    t0 = r["session_id"].map(rows.set_index("session_id")["t0"])
    e = r["session_id"].map(rows.set_index("session_id")["kwh_delivered"])
    # a quarter-hour is full when it starts at or after plug-in and the session's energy is not yet spent after it
    r = r.assign(cum=r.groupby("session_id")["kwh"].cumsum())
    full = (r["q"] >= t0) & (r["cum"] < e - 1e-9)
    m = 4 * r[full].groupby("session_id")["kwh"].max()
    return rows["session_id"].map(m).round(3)


def session_model(P, rec):
    """Largest gap between a shipped reading and the reading the constant-draw model gives it."""
    sp = P["sp"]
    r = sp[sp["session_id"].isin(rec["session_id"])]
    key = rec.set_index("session_id")
    r = r[r["version"] == r["session_id"].map(key["version"])]
    obs = observed(P, rec).fillna(0)
    rate = r["session_id"].map(pd.Series(obs.to_numpy(), index=rec["session_id"].to_numpy()))
    t0 = r["session_id"].map(key["t0"]).to_numpy(float)
    e = r["session_id"].map(key["kwh_delivered"]).to_numpy(float)
    rt = rate.to_numpy(float)
    ok = rt > 0
    end = t0 + np.where(ok, e / np.where(ok, rt, 1) * 3600, 0)
    q = r["q"].to_numpy(float)
    model = np.where(ok, rt * np.clip(np.minimum(q + Q, end) - np.maximum(q, t0), 0, None) / 3600, np.nan)
    gap = np.abs(model - r["kwh"].to_numpy())
    return float(np.nanmax(gap))


def vehicles_at(P, when):
    ch = P["checks"].copy()
    ch["checked_on"] = pd.to_datetime(ch["checked_on"]).dt.date
    ch = ch[ch["checked_on"] <= when].sort_values("checked_on").groupby("permit_no").last()
    per = P["permits"].set_index("permit_no")
    act = per[per["status"] == "ACTIVE"].index
    ch = ch.loc[ch.index.isin(act)].copy()
    ch["deck"] = per.loc[ch.index, "deck"]
    ref = P["ref"]
    ch["rating"] = [float(ref[(ref["make"].str.upper() == mk) & (ref["model"].str.upper() == md)
                              & (ref["model_year_from"] <= yr) & (ref["model_year_to"] >= yr)
                              & ((ref["trim"] == "") | (ref["trim"].str.upper() == tr))]["onboard_charger_kw"].iloc[0])
                    for mk, md, tr, yr in zip(ch["make"], ch["model"], ch["trim"], ch["model_year"])]
    return ch


def backtest_load(P, G, rec):
    """The decks' 2024 and 2025 load on the old units: the settlement export's sessions of record at their dated
    units, the gateway B sessions, and the fleet card charges the export does not carry."""
    decks = ["Civic Center North Deck", "Civic Center South Deck"]
    sp = P["sp"]
    deck_rec = rec[rec["garage"].isin(decks)]
    rr = sp[sp["session_id"].isin(deck_rec["session_id"])]
    rr = rr[rr["version"] == rr["session_id"].map(deck_rec.set_index("session_id")["version"])]
    load = G.readings_load(rr["q"], rr["kwh"])
    return load + other_loads(P, G)


def other_loads(P, G):
    """Gateway B's sessions and the fleet card charges the export does not carry, on the old units (computed once)."""
    if "_other" in P:
        return P["_other"]
    decks = ["Civic Center North Deck", "Civic Center South Deck"]
    load = np.zeros(G.n)
    gw = P["gw"].copy()
    gw["t0"] = [int(datetime.strptime(x, "%m/%d/%Y %H:%M:%S").replace(tzinfo=LA).timestamp()) for x in gw["Start"]]
    gw["d"] = [datetime.strptime(x, "%m/%d/%Y %H:%M:%S").date() for x in gw["Start"]]
    gw["station_id"] = gw["Station"]
    g_gar, _ = dated_join(P, gw)
    gw = gw[g_gar.isin(decks)]
    load = load + G.blocks(gw["t0"], gw["Energy (Wh)"] / 1000.0, np.full(len(gw), 6.6))
    fc = P["fc"].copy()
    fc["t0"] = [int(datetime.strptime(x, "%Y-%m-%d %H:%M:%S").replace(tzinfo=timezone.utc).timestamp())
                for x in fc["START"]]
    both = fc[fc["NETWORK_REF"].isin(set(P["hdr"]["auth_code"]))]
    t_exp = P["hdr"].drop_duplicates("auth_code").set_index("auth_code")["t0"]
    check("V19b the fleet card file stamps START in UTC: on every charge both files carry, START is the export's start "
          "to the second", len(both) > 1000 and (both["t0"].to_numpy() == both["NETWORK_REF"].map(t_exp).to_numpy()).all(),
          len(both))
    fc = fc[~fc["NETWORK_REF"].isin(set(P["hdr"]["auth_code"]))].copy()
    fc["d"] = [datetime.fromtimestamp(int(t), LA).date() for t in fc["t0"]]
    fc["station_id"] = fc["STATION"]
    f_gar, _ = dated_join(P, fc)
    fd = fc[f_gar.isin(decks)]
    load = load + G.blocks(fd["t0"], fd["KWH"].astype(float), np.full(len(fd), 6.6))
    P["_fleet_outside"] = (fc, fd)
    P["_other"] = load
    return load


def forecast_window(P):
    """FES-07: a forecast starts from the latest twelve closed calendar months, and a month is closed once its
    sessions have settled. A forecast from the 2024 base months can be made from the day the last December 2024 deck
    record settled to the day before the last January 2025 deck record settled."""
    h = P["hdr"].copy()
    h["garage"], _ = dated_join(P, h)
    h = h[h["garage"].isin(["Civic Center North Deck", "Civic Center South Deck"]) & (h["version"] == 1)]
    so = pd.to_datetime(h["settled_on"]).dt.date
    dec_close = max(so[(h["d"] >= date(2024, 12, 1)) & (h["d"] <= date(2024, 12, 31))])
    jan_close = max(so[(h["d"] >= date(2025, 1, 1)) & (h["d"] <= date(2025, 1, 31))])
    return dec_close, jan_close - timedelta(days=1)


def b3(P, R, G, rec):
    lo, hi = forecast_window(P)
    decks = ["Civic Center North Deck", "Civic Center South Deck"]
    bases = {}
    for when in (lo, hi):
        r0 = of_record_asof(P, when).copy()
        r0["garage"], r0["position"] = dated_join(P, r0)
        bases[when] = backtest_load(P, G, r0)
    load24 = bases[hi]
    load = backtest_load(P, G, rec)
    fc, fd = P["_fleet_outside"]
    dec = P["dec"].copy()
    dec["recv"] = pd.to_datetime(dec["received_on"]).dt.date
    late = dec[dec["recv"] > hi]
    late_24 = P["hdr"][P["hdr"]["session_id"].isin(late["session_id"]) & (P["hdr"]["d"] < date(2025, 1, 1))]
    check("V20b a forecast from the 2024 base can be made from %s to %s; twelve 2024 deck sessions were restated after "
          "that (received 3 to 20 March 2025, accepted), and the base as it stood is the same on every day of the window"
          % (lo, hi), lo <= hi and late_24["session_id"].nunique() == 12
          and all(near(G.peak(bases[lo], year=2024, month=m)[0], G.peak(bases[hi], year=2024, month=m)[0], 1e-9)
                  for m in range(1, 13)), (lo, hi, late_24["session_id"].nunique()))
    check("V19a the fleet card charges outside the export all predate the April 2025 platform move",
          len(fd) > 0 and max(fc["d"]) < date(2025, 4, 1), (len(fd), max(fc["d"])))
    fac = factor_on(R["factors"], date(2025, 1, 31))
    check("V20 the factor in force for a forecast made from 2024 is 1.08", fac == 1.08, fac)
    st = re.sub(r"\s+", " ", pdf_text(os.path.join(os.path.dirname(P["log_path"]), "fes-07_load_forecasting_standard_rev4.pdf")))
    check("V19 the standard files the accuracy record: whole-kW forecast less whole-kW recorded demand",
          "Forecast demand is stated in whole kilowatts" in st and "recorded billing demand in whole kilowatts" in st)
    F = {m: fac * G.peak(load24, year=2024, month=m)[0] for m in range(1, 13)}
    A = {m: G.peak(load, year=2025, month=m)[0] for m in range(1, 13)}
    Fq = {m: int(math.floor(fac * G.peak(load, year=2024, month=m)[0] + 0.5)) for m in range(1, 13)}
    check("V20c the 2024 base taken on the restated versions moves all twelve forecasts",
          all(Fq[m] != int(math.floor(F[m] + 0.5)) for m in range(1, 13)), Fq)
    fc = {m: int(math.floor(F[m] + 0.5)) for m in F}
    ac = {m: int(math.floor(A[m] + 0.5)) for m in A}
    miss = {m: 100 * (fc[m] - ac[m]) / ac[m] for m in F}
    ms = {m: round(miss[m], 1) for m in F}
    check("V21 back-test misses within +-3.0 per cent, at least four each way",
          max(abs(x) for x in miss.values()) <= 3.0 and sum(x > 0 for x in miss.values()) >= 4
          and sum(x < 0 for x in miss.values()) >= 4, ms)
    check("V22 every forecast and recorded demand sits off a whole kW and inside its bin; no miss on a bin edge",
          all(0.06 <= x % 1 <= 0.24 or 0.76 <= x % 1 <= 0.94 for x in list(F.values()) + list(A.values()))
          and all(0.004 <= abs(v - round(v, 1)) <= 0.035 for v in miss.values()))
    print("     B3 forecasts", fc)
    print("     B3 misses   ", ms)
    check("V22a the 24 back-test figures match the key", [fc[m] for m in range(1, 13)] == EXPECTED["b3_forecast"]
          and [ms[m] for m in range(1, 13)] == EXPECTED["b3_miss"])
    return {"forecast": fc, "miss": ms}


PST = timezone(timedelta(hours=-8))


def panel_frames(P, rec):
    sched = P["sched"].copy()
    sched["to"] = sched["effective_to"].fillna(pd.Timestamp("2099-12-31"))
    ev = sched[sched["description"].str.startswith("EVSE ")].copy()
    ev["pos"] = ev["description"].str.extract(r"EVSE ([NS]-\d\d)")[0]
    sp = P["sp"]
    deck = rec[rec["position"].fillna("").str.match(r"^[NS]-\d\d$") & rec["garage"].isin(
        ["Civic Center North Deck", "Civic Center South Deck"])].copy()
    deck["day"] = deck["d"]
    deck["end"] = deck["t0"] + deck["kwh_delivered"] / 6.6 * 3600.0
    r = sp[sp["session_id"].isin(deck["session_id"])]
    r = r[r["version"] == r["session_id"].map(deck.set_index("session_id")["version"])]
    r = r.assign(pos=r["session_id"].map(deck.set_index("session_id")["position"]))
    r["qd"] = pd.to_datetime(r["q"], unit="s", utc=True).dt.tz_convert(LA).dt.date
    return ev, deck, r


def panel_spans(P, G, rec):
    """Each 2026 read span of each panel meter, on the meter's own clock (standard time all year), the later of two
    reads on a day standing, against the session energy through the panel between the two read instants: readings
    for every whole quarter-hour, and the quarter-hour a read falls inside on the sessions' constant 6.6 kW draw."""
    log = P["log"]
    ev, deck, r = panel_frames(P, rec)
    out = {}
    for panel in ("CP-N", "CP-S"):
        evp = ev[ev["panel"] == panel]
        rows = r.merge(evp[["pos", "effective_from", "to"]], on="pos")
        rows = rows[(rows["effective_from"].dt.date <= rows["qd"]) & (rows["to"].dt.date >= rows["qd"])]
        ss = deck.merge(evp[["pos", "effective_from", "to"]], left_on="position", right_on="pos")
        ss = ss[(ss["effective_from"].dt.date <= ss["day"]) & (ss["to"].dt.date >= ss["day"])]
        qs = rows.groupby("q")["kwh"].sum().sort_index()
        cq = np.concatenate([[0.0], np.cumsum(qs.to_numpy())])
        qi = qs.index.to_numpy()
        st, en = ss["t0"].to_numpy(float), ss["end"].to_numpy(float)

        def energy_to(t):
            q0 = G.t0 + Q * ((t - G.t0) // Q)
            whole = cq[np.searchsorted(qi, q0)]
            lo, hi = np.maximum(st, q0), np.minimum(en, t)
            return whole + float(np.sum(np.where(hi > lo, 6.6 * (hi - lo), 0.0)) / 3600.0)
        cts = P["ct"].copy()
        cts["garage"], cts["position"] = dated_join(P, cts)
        cts = cts.merge(evp[["pos", "effective_from", "to"]], left_on="position", right_on="pos")
        cts = cts[(cts["effective_from"].dt.date <= cts["d"]) & (cts["to"].dt.date >= cts["d"])]
        ct_t0, ct_t1, ct_e = cts["t0"].to_numpy(), cts["t1"].to_numpy(), cts["Energy (kWh)"].to_numpy(float)

        def courtesy_in(a, b):
            assert not ((ct_t0 < a) & (ct_t1 > a)).any() and not ((ct_t0 < b) & (ct_t1 > b)).any()
            return float(ct_e[(ct_t0 >= a) & (ct_t0 < b)].sum())
        lg = log[log["Panel"] == panel].copy()
        lg = lg.groupby("Read date").last().reset_index()
        # a corrected reading replaces the reading logged for that date
        fx = P["fixes"][P["fixes"]["Panel"] == panel].set_index("Read date")
        for i in lg.index:
            d = lg.at[i, "Read date"]
            if d in fx.index:
                assert abs(float(lg.at[i, "kWh register"]) - float(fx.at[d, "kWh as logged"])) < 1e-9
                lg.at[i, "kWh register"] = float(fx.at[d, "kWh corrected"])
        lg["t"] = [int(datetime.combine(d, datetime.strptime(tm, "%H:%M").time(), tzinfo=PST).timestamp())
                   for d, tm in zip(lg["Read date"], lg["Read time (meter)"])]
        lg = lg.sort_values("t")
        lg = lg[lg["Read date"] >= date(2025, 12, 31)].reset_index(drop=True)
        spans = []
        for i in range(1, len(lg)):
            a, b = lg.loc[i - 1], lg.loc[i]
            sess = energy_to(b["t"]) - energy_to(a["t"]) + courtesy_in(a["t"], b["t"])
            metered = float(b["kWh register"] - a["kWh register"])
            spans.append({"read": i, "meter": b["Meter"], "date": b["Read date"], "time": b["Read time (meter)"],
                          "start": datetime.fromtimestamp(int(a["t"]), LA), "end": datetime.fromtimestamp(int(b["t"]), LA),
                          "metered": metered, "sessions": sess, "gap": metered - sess})
        out[panel] = spans
    return out, ev, r


def b1(P, R, G, rec):
    log = P["log"]
    plates = P["plates"]
    check("V23 the deck sub-meters keep standard time all year",
          all("Standard Time" in x for x in plates["clock_time_base"]) and set(plates["dst_adjustment"]) == {"Disabled"})
    check("V26 every read is logged at the close of a quarter-hour on the meter clock",
          all(int(x.split(":")[1]) % 15 == 0 for x in log["Read time (meter)"]))
    fx = P["fixes"]
    check("V27 the log's Corrections sheet corrects eight 2026 reads, four on each meter, each matching the reading "
          "logged on the Reads sheet", len(fx) == 8 and fx.groupby("Panel").size().to_dict() == {"CP-N": 4, "CP-S": 4}
          and all(((log["Read date"] == d) & (log["Panel"] == pn) & ((log["kWh register"] - k).abs() < 1e-9)).any()
                  for d, pn, k in zip(fx["Read date"], fx["Panel"], fx["kWh as logged"])), len(fx))
    spans, ev, r = panel_spans(P, G, rec)
    out = {p: [x["gap"] for x in v] for p, v in spans.items()}
    for panel in ("CP-N", "CP-S"):
        print(f"     B1 {panel}", [int(math.floor(v + 0.5)) for v in out[panel]])
    pst = PST
    bad = []
    for panel in ("CP-N", "CP-S"):
        evp = ev[ev["panel"] == panel]
        rows = r.merge(evp[["pos", "effective_from", "to"]], on="pos")
        rows = rows[(rows["effective_from"].dt.date <= rows["qd"]) & (rows["to"].dt.date >= rows["qd"])]
        q = rows.groupby("q")["kwh"].sum() * 4
        lg = log[log["Panel"] == panel].copy()
        lg["t"] = [int(datetime.combine(d, datetime.strptime(tm, "%H:%M").time(), tzinfo=pst).timestamp())
                   for d, tm in zip(lg["Read date"], lg["Read time (meter)"])]
        lg = lg.sort_values("t").reset_index(drop=True)
        for i in range(1, len(lg)):
            if lg.loc[i, "Read date"] < date(2026, 1, 1):
                continue
            seg = q[(q.index >= lg.loc[i - 1, "t"]) & (q.index < lg.loc[i, "t"])]
            if abs(round(float(seg.max()), 1) - float(lg.loc[i, "Max demand (kW)"])) > 1e-6:
                bad.append((panel, str(lg.loc[i, "Read date"]), float(seg.max()), lg.loc[i, "Max demand (kW)"]))
    check("V25 every 2026 maximum-demand register reproduces from the session readings", not bad, bad)
    check("V24a the 24 panel figures match the key", {p: [int(math.floor(x + 0.5)) for x in v] for p, v in out.items()}
          == EXPECTED["b1"])
    check("V24 B1: 24 readings, each a plausible unaccounted load 0.06 to 0.24 kWh off a whole kWh",
          all(len(v) == 12 for v in out.values()) and all(0 < x < 1500 and (0.06 <= x % 1 <= 0.24 or 0.76 <= x % 1 <= 0.94)
                                                           for v in out.values() for x in v))
    return {p: [int(math.floor(x + 0.5)) for x in v] for p, v in out.items()}


if __name__ == "__main__":
    main()
