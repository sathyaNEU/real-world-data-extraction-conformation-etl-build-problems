"""The pack's documents: the utility's rate schedule and planning guide, the city's forecasting
standard, the draft service agreement, the network's export field notes and Parking Services'
data-sources note. Short, functional, written in each organisation's own voice."""
from __future__ import annotations

import re
import zipfile
from datetime import date, datetime, timezone
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

from reportlab import rl_config

rl_config.invariant = 1  # fixed document identifier and dates, so two builds are byte-identical

from reportlab.lib import colors  # noqa: E402
from reportlab.lib.enums import TA_LEFT  # noqa: E402
from reportlab.lib.pagesizes import letter  # noqa: E402
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # noqa: E402
from reportlab.lib.units import inch  # noqa: E402
from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)  # noqa: E402

from common import tariff_holidays  # noqa: E402


def _styles():
    ss = getSampleStyleSheet()
    base = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=12.6,
                          alignment=TA_LEFT, spaceAfter=5)
    small = ParagraphStyle("s", parent=base, fontSize=8, leading=10, textColor=colors.HexColor("#333333"))
    h1 = ParagraphStyle("h1", parent=base, fontName="Helvetica-Bold", fontSize=12.5, leading=15, spaceAfter=3)
    h2 = ParagraphStyle("h2", parent=base, fontName="Helvetica-Bold", fontSize=10, leading=13, spaceBefore=6,
                        spaceAfter=3)
    return base, small, h1, h2


def _table(data, widths, header=True):
    t = Table(data, colWidths=widths, hAlign="LEFT")
    style = [("FONT", (0, 0), (-1, -1), "Helvetica", 8.5),
             ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.black),
             ("BOTTOMPADDING", (0, 0), (-1, -1), 2), ("TOPPADDING", (0, 0), (-1, -1), 2)]
    if header:
        style.append(("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.5))
    t.setStyle(TableStyle(style))
    return t


LA = ZoneInfo("America/Los_Angeles")
INVARIANT_PDF_DATE = b"D:20000101000000+00'00'"   # what reportlab writes for both dates under rl_config.invariant

# the local clock time each filed PDF was produced on its issue date
PDF_MADE = {"rates": datetime(2025, 11, 14, 10, 47, 23), "standard": datetime(2026, 9, 15, 15, 29, 6),
            "guide": datetime(2026, 3, 2, 9, 31, 18)}


def pdf_date(when: datetime) -> bytes:
    """A PDF date for a local civil time, with the Pacific offset in force that day."""
    off = int(when.replace(tzinfo=LA).utcoffset().total_seconds() // 60)
    sign, off = ("-" if off < 0 else "+"), abs(off)
    return f"D:{when:%Y%m%d%H%M%S}{sign}{off // 60:02d}'{off % 60:02d}'".encode()


def stamp_pdf(path, when: datetime):
    """Creation and modification dates set to the time the document was produced. Same byte length as the invariant
    date, so the cross-reference table stays valid."""
    new = pdf_date(when)
    assert len(new) == len(INVARIANT_PDF_DATE), new
    b = open(path, "rb").read()
    assert b.count(INVARIANT_PDF_DATE) == 2, path
    with open(path, "wb") as f:
        f.write(b.replace(INVARIANT_PDF_DATE, new))


def _doc(path, title, author, subject, story, footer, when=None):
    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.drawString(0.85 * inch, 0.55 * inch, footer)
        canvas.drawRightString(7.65 * inch, 0.55 * inch, "Page %d" % doc.page)
        canvas.restoreState()

    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.85 * inch, rightMargin=0.85 * inch,
                            topMargin=0.8 * inch, bottomMargin=0.85 * inch, title=title, author=author,
                            subject=subject, creator=author)
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    if when is not None:
        stamp_pdf(path, when)


def _ds(d: date) -> str:
    return f"{d.strftime('%B')} {d.day}, {d.year}"


# ------------------------------------------------------------------------------ rate schedule
def rate_schedule(path):
    base, small, h1, h2 = _styles()
    s = []
    s.append(Paragraph("NORTH SOUND POWER &amp; LIGHT COMPANY", small))
    s.append(Paragraph("Tariff Book, Electric Service", small))
    s.append(Spacer(1, 6))
    s.append(Paragraph("SCHEDULE 26<br/>ELECTRIC VEHICLE CHARGING SERVICE, SECONDARY VOLTAGE", h1))
    s.append(Paragraph("Original Sheet No. 26-1 &nbsp;&nbsp; Issued November 14, 2025 &nbsp;&nbsp; "
                       "Effective with service on and after January 1, 2026", small))
    s.append(Spacer(1, 8))
    s.append(Paragraph("1. Availability", h2))
    s.append(Paragraph(
        "Available to nonresidential Customers for a separately metered secondary-voltage service that supplies "
        "electric vehicle charging equipment and no other load. Service is single-phase or three-phase at the "
        "Company's option, through one interval meter per service.", base))
    s.append(Paragraph("2. Monthly Rate", h2))
    s.append(_table([["Charge", "Rate"],
                     ["Basic Charge", "$185.00 per month"],
                     ["Contract Demand Charge", "$11.40 per kW of Contract Demand per month"],
                     ["Energy Charge, On-Peak", "13.47 cents per kWh"],
                     ["Energy Charge, Off-Peak", "9.12 cents per kWh"]], [2.6 * inch, 3.6 * inch]))
    s.append(Spacer(1, 4))
    s.append(Paragraph(
        "On-Peak energy is energy delivered in the intervals counted for Billing Demand below. All other energy is "
        "Off-Peak.", base))
    s.append(Paragraph("3. Billing Demand", h2))
    s.append(Paragraph(
        "The Billing Demand for a month is the highest average kilowatt load measured in any fifteen-minute "
        "interval beginning at 12:00 noon through 7:45 p.m., Monday through Friday, excluding the holidays listed "
        "in Section 7. For interval-metered service the billing month is the calendar month.", base))
    s.append(Paragraph("4. Contract Demand", h2))
    s.append(Paragraph(
        "The Customer shall contract for a Contract Demand stated in whole multiples of 5 kW for each Contract "
        "Year. The Contract Demand Charge applies to the Contract Demand in every month of the Contract Year, "
        "whatever the Billing Demand of the month.", base))
    s.append(Paragraph(
        "If the Billing Demand in any month exceeds the Contract Demand, the Contract Demand is reset to that "
        "Billing Demand, rounded up to the next multiple of 5 kW, for that month and the following eleven months.",
        base))
    s.append(Paragraph("5. Metering", h2))
    s.append(Paragraph(
        "The Company will install a fifteen-minute interval meter. Interval data are stored in local prevailing "
        "time and are available to the Customer through the Company's account portal.", base))
    s.append(Paragraph("6. Term", h2))
    s.append(Paragraph(
        "Service under this schedule requires a written service agreement with an initial term of not less than "
        "three Contract Years.", base))
    s.append(Paragraph("7. Holidays", h2))
    s.append(Paragraph(
        "New Year's Day, Memorial Day, Independence Day, Labor Day, Thanksgiving Day and Christmas Day. A holiday "
        "falling on a Saturday is observed on the preceding Friday; a holiday falling on a Sunday is observed on the "
        "following Monday. Observed dates:", base))
    rows = [["Holiday", "2024", "2025", "2026", "2027", "2028"]]
    hol = {y: dict(tariff_holidays(y)) for y in range(2024, 2029)}
    for name in ["New Year's Day", "Memorial Day", "Independence Day", "Labor Day", "Thanksgiving Day",
                 "Christmas Day"]:
        rows.append([name] + [hol[y][name].strftime("%b %-d") + ("" if hol[y][name].year == y else
                                                                 f", {hol[y][name].year}") for y in range(2024, 2029)])
    s.append(_table(rows, [1.5 * inch] + [0.95 * inch] * 5))
    s.append(Spacer(1, 6))
    s.append(Paragraph("8. General Rules", h2))
    s.append(Paragraph(
        "Service under this schedule is subject to the Company's General Rules and Regulations and to the "
        "Company's Rule 16, Line and Service Extensions.", base))
    _doc(path, "Schedule 26 Electric Vehicle Charging Service", "North Sound Power & Light", "Tariff Book",
         s, "NSPL Tariff Book  |  Schedule 26  |  Issued November 14, 2025", when=PDF_MADE["rates"])


# ------------------------------------------------------------------------------ forecasting standard
FACTORS = [("2024", "1.06", date(2023, 9, 12), "County EV registrations, Q2 2023 against Q2 2022"),
           ("2025", "1.08", date(2024, 9, 10), "County EV registrations, Q2 2024 against Q2 2023"),
           ("2025", "1.09", date(2025, 4, 8), "Mid-year revision, Q1 2025 registration count"),
           ("2026", "1.10", date(2025, 9, 9), "County EV registrations, Q2 2025 against Q2 2024"),
           ("2027", "1.12", date(2026, 9, 15), "County EV registrations, Q2 2026 against Q2 2025")]


def forecasting_standard(path):
    base, small, h1, h2 = _styles()
    s = []
    s.append(Paragraph("CITY OF LARCH HARBOR &nbsp;|&nbsp; PUBLIC WORKS, ENERGY &amp; FACILITIES DIVISION", small))
    s.append(Spacer(1, 4))
    s.append(Paragraph("Facilities Electrical Load Forecasting Standard<br/>FES-07, Revision 4", h1))
    s.append(Paragraph("Issued September 15, 2026 &nbsp;&nbsp; Owner: City Energy Manager &nbsp;&nbsp; "
                       "Supersedes Revision 3 (September 9, 2025)", small))
    s.append(Spacer(1, 8))
    s.append(Paragraph("1. Scope", h2))
    s.append(Paragraph(
        "This standard applies to every load forecast the City files with an electric utility for a new or "
        "altered service, and to forecasts that support a capacity request to City Council. It does not apply "
        "to building energy budgets, which follow FES-03.", base))
    s.append(Paragraph("2. Base months", h2))
    s.append(Paragraph(
        "A forecast for a new or altered service starts from the latest twelve closed calendar months of metered "
        "or settled demand at the equipment the service will supply. Each forecast month is built from the same "
        "calendar month of the base year. A month is closed once its sessions or meter reads have settled.", base))
    s.append(Paragraph(
        "Where the equipment the service will supply differs from the equipment in service during the base "
        "months, the forecast is prepared for the equipment the service will supply.", base))
    s.append(Paragraph("3. Growth", h2))
    s.append(Paragraph(
        "The county EV registration growth factor in Table 1 is applied to each forecast month's demand and "
        "energy. A factor applies to forecasts made on or after its adoption date; a forecast uses the factor in "
        "force on the date it is made.", base))
    rows = [["Planning year", "Factor", "Adopted", "Basis"]]
    for py, f, d, b in FACTORS:
        rows.append([py, f, _ds(d), b])
    s.append(Paragraph("Table 1. County EV registration growth factor", small))
    s.append(_table(rows, [1.0 * inch, 0.7 * inch, 1.5 * inch, 3.6 * inch]))
    s.append(Spacer(1, 6))
    s.append(Paragraph("4. Records and accuracy", h2))
    s.append(Paragraph(
        "Forecast demand is stated in whole kilowatts. The base-month demand is carried unrounded; the forecast is "
        "rounded once, after the factor is applied. The forecast workbook, the base-month records as they stood on "
        "the day the forecast was made, and the factor used are kept with the filing for six years; a forecast is not "
        "restated when a base-month record is later restated. Forecasts presented to Council state the base months "
        "and the factor.", base))
    s.append(Paragraph(
        "Once a forecast month has closed, the division records the forecast's error: the stated forecast less "
        "the month's recorded billing demand in whole kilowatts, as a percentage of that recorded demand.", base))
    s.append(Paragraph("5. Revision history", h2))
    s.append(_table([["Revision", "Date", "Change"],
                     ["2", "September 2023", "Factor table moved into this standard from FES-03"],
                     ["3", "September 2025", "Base months set at the equipment the service will supply"],
                     ["4", "September 2026", "2027 factor added"]], [0.8 * inch, 1.4 * inch, 4.6 * inch]))
    _doc(path, "FES-07 Facilities Electrical Load Forecasting Standard", "City of Larch Harbor",
         "Energy & Facilities Division", s, "City of Larch Harbor  |  FES-07 Rev. 4  |  Internal standard",
         when=PDF_MADE["standard"])


# ------------------------------------------------------------------------------ planning guide
def planning_guide(path):
    base, small, h1, h2 = _styles()
    s = []
    s.append(Paragraph("NORTH SOUND POWER &amp; LIGHT COMPANY &nbsp;|&nbsp; DISTRIBUTION PLANNING", small))
    s.append(Spacer(1, 4))
    s.append(Paragraph("New Service Planning Guide, 2026 Edition<br/>Section 7: Electric Vehicle Charging Loads",
                       h1))
    s.append(Paragraph("Edition date March 2, 2026 &nbsp;&nbsp; For customers, contractors and NSPL planners",
                       small))
    s.append(Spacer(1, 8))
    s.append(Paragraph("7.1 Application", h2))
    s.append(Paragraph(
        "Customers requesting a new service for EV charging equipment submit the equipment list with nameplate "
        "ratings and the planned in-service date. Requests for more than 100 kW of connected charging load are "
        "assigned a planner, who schedules a service review with the customer.", base))
    s.append(Paragraph("7.2 Information the planner needs", h2))
    s.append(Paragraph(
        "Unit count, unit nameplate rating, the panel and feeder arrangement, any load management controls, and "
        "the customer's operating hours. Where existing equipment is being replaced, the customer supplies twelve "
        "months of metered or settled usage for it if available.", base))
    s.append(Paragraph("7.3 Sizing", h2))
    s.append(Paragraph(
        "NSPL planners size a new EV charging service on the connected nameplate rating of the charging units "
        "multiplied by the diversity factor in Table 7-2, stated to the next 5 kW above. The same sizing sets the "
        "service transformer and the secondary conductors.", base))
    s.append(Paragraph("Table 7-2. Diversity factors for EV charging units", small))
    s.append(_table([["Units on the service", "Diversity factor"], ["1 to 5", "1.00"], ["6 to 20", "0.75"],
                     ["21 to 40", "0.60"], ["41 to 80", "0.50"], ["Over 80", "0.45"]], [2.2 * inch, 1.6 * inch]))
    s.append(Spacer(1, 6))
    s.append(Paragraph("7.4 Service review", h2))
    s.append(Paragraph(
        "The planner presents the sizing at the service review, with the transformer and conductor selection "
        "and the construction schedule. The review is normally held six to ten weeks before energization.", base))
    s.append(Paragraph("7.5 Construction", h2))
    s.append(Paragraph(
        "The customer's electrician installs the service equipment and the customer-side distribution. NSPL "
        "sets the meter after the jurisdiction's electrical inspection is approved.", base))
    _doc(path, "New Service Planning Guide, Section 7", "North Sound Power & Light", "Distribution Planning", s,
         "NSPL New Service Planning Guide  |  2026 Edition  |  Section 7", when=PDF_MADE["guide"])


# ------------------------------------------------------------------------------ service agreement (docx)
def service_agreement(path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)

    def p(text, bold=False, size=None, align=None, space=4):
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if align:
            para.alignment = align
        para.paragraph_format.space_after = Pt(space)
        return para

    p("DRAFT for Council packet, February 16, 2027", size=9)
    p("ELECTRIC SERVICE AGREEMENT", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space=2)
    p("New Service, Schedule 26", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space=10)
    p("This Agreement is between North Sound Power & Light Company (\"NSPL\") and the City of Larch Harbor, a "
      "Washington municipal corporation (the \"Customer\"). Draft dated January 12, 2027.")
    p("1. Premises and service", bold=True)
    p("NSPL will provide a new secondary-voltage service under Schedule 26 to the Civic Center North Deck and "
      "Civic Center South Deck, 600 block of Harbor Avenue. The new service replaces the decks' existing feed "
      "from the Civic Center campus service.")
    p("2. Load served", bold=True)
    p("The service supplies the thirty-two (32) electric vehicle charging units listed in Exhibit A and no other "
      "load. Both deck charging panels are fed from the new service through one NSPL revenue meter.")
    p("3. Energization", bold=True)
    p("NSPL will energize the service on or about April 1, 2027, after the Customer's contractor removes the "
      "existing charging units and installs the units in Exhibit A.")
    p("4. Contract Year and Contract Demand", bold=True)
    p("The first Contract Year is the twelve billing months beginning April 1, 2027. The Customer states in "
      "Schedule 1 the maximum Billing Demand it expects in the first Contract "
      "Year, and that figure is the Contract Demand for the first Contract Year. The Customer shall return "
      "Schedule 1 to NSPL no later than March 1, 2027. Contract Demand for later Contract Years is set under "
      "Schedule 26.")
    p("5. Term", bold=True)
    p("The initial term is three Contract Years from energization and continues year to year after that until "
      "either party gives ninety days' written notice.")
    p("6. Approval", bold=True)
    p("This Agreement takes effect when approved by the Larch Harbor City Council and signed by both parties.")
    p("7. Notices", bold=True)
    p("NSPL: Paul Henderson, New Service Planning, North Sound Power & Light Company. Customer: Shelley Tanner, "
      "City Energy Manager, Public Works, Energy & Facilities Division.", space=12)
    p("EXHIBIT A. CHARGING UNITS SUPPLIED BY THE SERVICE", bold=True)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for c, t in zip(hdr, ["Location", "Units", "Unit", "Rating"]):
        c.text = t
    for loc in ("Civic Center North Deck, levels 2 and 3", "Civic Center South Deck, levels 2 and 3"):
        cells = table.add_row().cells
        for c, t in zip(cells, [loc, "16", "Single-port Level 2, 48 A continuous at 240 V", "11.5 kW"]):
            c.text = t
    p("")
    p("The Exhibit A units replace thirty-two single-port units rated 6.6 kW, which the Customer will remove "
      "before energization.", space=12)
    p("SCHEDULE 1. CONTRACT DEMAND, FIRST CONTRACT YEAR", bold=True)
    p("Contract Demand for the first Contract Year (whole multiples of 5 kW): ____________ kW")
    p("Signed for the Customer: ______________________   Date: ____________", space=14)
    title = "Electric Service Agreement, New Service, Schedule 26"
    doc.core_properties.title = title
    doc.core_properties.author = "North Sound Power & Light"
    doc.save(path)
    texts = [x.text for x in doc.paragraphs]
    cells = [[c.text for c in row.cells] for t in doc.tables for row in t.rows]
    _finish_docx(path, texts, cells, title)


# NSPL's draft as last saved: started January 8, last saved 3:40 p.m. on January 12, 2027 (the draft date)
DOCX_SAVED = {"created": datetime(2027, 1, 8, 17, 22, tzinfo=timezone.utc),
              "modified": datetime(2027, 1, 12, 23, 40, tzinfo=timezone.utc),
              "by": "Paul Henderson", "revision": 6, "minutes": 74}
# characters per line of body text and of an Exhibit A cell at Calibri 10.5 on the template's 6-inch text width, and
# lines per page once the headings' spacing is allowed for (a render of the draft runs onto a second page)
DOCX_LINE_CHARS, DOCX_CELL_CHARS, DOCX_PAGE_LINES = 92, 22, 40


def docx_stats(texts, cells):
    """Word's document statistics for the text: words, characters with and without spaces, paragraphs, lines."""
    flat = [t for t in texts if t.strip()] + [c for row in cells for c in row if c.strip()]
    words = sum(len(t.split()) for t in flat)
    chars = sum(len(re.sub(r"\s", "", t)) for t in flat)
    spaced = sum(len(t) for t in flat)
    lines = sum(-(-len(t) // DOCX_LINE_CHARS) for t in texts if t.strip())
    lines += sum(max(-(-len(c) // DOCX_CELL_CHARS) for c in row) for row in cells)
    return {"words": words, "chars": chars, "spaced": spaced, "paragraphs": len(flat), "lines": lines,
            "pages": 1 + (lines - 1) // DOCX_PAGE_LINES}


def _finish_docx(path, texts, cells, title):
    """Package properties as Word writes them on save (document statistics, author, revision, save times), and no
    preview picture: the template's own thumbnail is a blank page."""
    st = docx_stats(texts, cells)
    iso = lambda d: d.strftime("%Y-%m-%dT%H:%M:%SZ")  # noqa: E731
    org = escape("North Sound Power & Light")
    core = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
            '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
            'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
            'xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
            f'<dc:title>{escape(title)}</dc:title><dc:subject></dc:subject><dc:creator>{org}</dc:creator>'
            f'<cp:keywords></cp:keywords><dc:description></dc:description>'
            f'<cp:lastModifiedBy>{DOCX_SAVED["by"]}</cp:lastModifiedBy><cp:revision>{DOCX_SAVED["revision"]}</cp:revision>'
            f'<dcterms:created xsi:type="dcterms:W3CDTF">{iso(DOCX_SAVED["created"])}</dcterms:created>'
            f'<dcterms:modified xsi:type="dcterms:W3CDTF">{iso(DOCX_SAVED["modified"])}</dcterms:modified>'
            '</cp:coreProperties>')
    app = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
           '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
           'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
           f'<Template>Normal.dotm</Template><TotalTime>{DOCX_SAVED["minutes"]}</TotalTime>'
           f'<Pages>{st["pages"]}</Pages><Words>{st["words"]}</Words><Characters>{st["chars"]}</Characters>'
           '<Application>Microsoft Office Word</Application><DocSecurity>0</DocSecurity>'
           f'<Lines>{st["lines"]}</Lines><Paragraphs>{st["paragraphs"]}</Paragraphs><ScaleCrop>false</ScaleCrop>'
           '<HeadingPairs><vt:vector size="2" baseType="variant"><vt:variant><vt:lpstr>Title</vt:lpstr></vt:variant>'
           '<vt:variant><vt:i4>1</vt:i4></vt:variant></vt:vector></HeadingPairs><TitlesOfParts>'
           f'<vt:vector size="1" baseType="lpstr"><vt:lpstr>{escape(title)}</vt:lpstr></vt:vector></TitlesOfParts>'
           f'<Company>{org}</Company><LinksUpToDate>false</LinksUpToDate>'
           f'<CharactersWithSpaces>{st["spaced"]}</CharactersWithSpaces><SharedDoc>false</SharedDoc>'
           '<HyperlinksChanged>false</HyperlinksChanged><AppVersion>14.0000</AppVersion></Properties>')
    with zipfile.ZipFile(path) as z:
        infos = [i for i in z.infolist() if i.filename != "docProps/thumbnail.jpeg"]
        data = {i.filename: z.read(i.filename) for i in infos}
    data["docProps/core.xml"] = core.encode()
    data["docProps/app.xml"] = app.encode()
    rels = data["_rels/.rels"].decode()
    rels, n = re.subn(r'<Relationship Id="rId\d+" Type="[^"]*/metadata/thumbnail" Target="docProps/thumbnail\.jpeg"/>',
                      "", rels)
    assert n == 1, rels
    data["_rels/.rels"] = rels.encode()
    ct = data["[Content_Types].xml"].decode()
    ct, n = re.subn(r'<Default Extension="jpeg" ContentType="image/jpeg"/>', "", ct)
    assert n == 1, ct
    data["[Content_Types].xml"] = ct.encode()
    stg = data["word/settings.xml"].decode()
    stg, n = re.subn(r"<w:savePreviewPicture/>", "", stg)
    assert n == 1
    data["word/settings.xml"] = stg.encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for i in infos:
            zi = zipfile.ZipInfo(i.filename, date_time=i.date_time)
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, data[i.filename])
    return st


# ------------------------------------------------------------------------------ field notes (txt)
FIELD_NOTES = """Curbline Charging
Settlement export field notes, version 3.4 (April 2025)

These notes describe the two files of the settlement export. They apply to every delivery
made on the current platform and to the history migrated to it.

settled_sessions
  session_id      Identifier given to a session record when it is delivered.
  version         Settlement version of the session. A restated session is delivered again
                  under the same session_id with the next version number.
  auth_code       Authorization code of the charge. One code identifies one settled charge;
                  a record delivered again keeps its code.
  station_id      Station identifier in use at the time of the session. Identifiers are drawn
                  from the network pool, and an identifier freed when a station is retired may
                  be assigned to another station. The station register lists each assignment
                  with its in-service dates.
  plug_in         Start of the session, local civil time with UTC offset.
  plug_out        End of the session, local civil time with UTC offset.
  kwh_delivered   Energy delivered to the vehicle in the session, metered at the station.
  account_type    PERMIT, FLEET or PUBLIC.
  permit_no       Parking permit presented at the station (PERMIT sessions).
  fleet_card      Fleet card presented at the station (FLEET sessions).
  settled_on      Date of the settlement run that settled the session. The settlement run
                  is daily, starting at 10:00 a.m. local time.
  delivered_on    Date of the weekly delivery that carried the row.

session_intervals
  session_id      As above.
  version         As above.
  interval_start  Start of the quarter-hour, local civil time (America/Los_Angeles).
  kwh             Energy delivered in the quarter-hour, metered at the station. Rows run from
                  the quarter-hour of plug_in through the quarter-hour of plug_out. A row of
                  zero is a quarter-hour in which the vehicle was connected and no energy was
                  delivered.

Average demand in a quarter-hour, in kW, is four times the quarter-hour's kWh.

Platform migration
  On April 1, 2025 every station was given a new identifier on the current platform.
  Records from before that date carry the identifier then in use.

"""


def field_notes(path):
    with open(path, "w", newline="\n") as f:
        f.write(FIELD_NOTES)


def data_sources(path, files_rows):
    lines = ["City of Larch Harbor, Parking Services",
             "Civic Center charging service: folder contents and data sources",
             "Prepared January 19, 2027 by A. Warner for the City Energy Manager",
             "",
             "Both Civic Center decks are permit-only. Each EV charging permit names one vehicle, and",
             "Parking Services checks the plate against the vehicle's registration at issue and at",
             "every January renewal.",
             ""]
    for name, text in files_rows:
        lines.append(name)
        for chunk in _wrap(text, 92):
            lines.append("    " + chunk)
        lines.append("")
    with open(path, "w", newline="\n") as f:
        f.write("\n".join(lines).rstrip() + "\n")


def _wrap(text, width):
    words, out, cur = text.split(), [], ""
    for w in words:
        if len(cur) + len(w) + 1 > width:
            out.append(cur)
            cur = w
        else:
            cur = (cur + " " + w).strip()
    if cur:
        out.append(cur)
    return out
