"""Civic Center decks, first contract year: the Council note, the load-day chart and the workbook.

    python3 task117/generator/golden.py [target_dir] [out_dir]

Reads only the files in target_dir (default ../target). Rebuilds the 2026 deck sessions, the car
behind each one and the replay on the Exhibit A units, then writes contract_demand_note.pdf,
deck_load_day.png and civic_service_demand.xlsx to out_dir (default ../golden) and prints the
figures the note rests on. The file readers and quarter-hour arithmetic are verify_pack.py's, which
reads nothing but the shipped files, so the goldens and the verifier cannot drift apart.
"""
from __future__ import annotations

import math
import os
import re
import subprocess
import sys
import zipfile
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from verify_pack import (EXPECTED, LA, Q, Grid, backtest_load, car_ratings, charges, dated_join,  # noqa: E402
                         factor_on, load, of_record, panel_spans, rules)

import matplotlib  # noqa: E402
matplotlib.use("Agg")
import matplotlib.dates as mdates  # noqa: E402
import matplotlib.pyplot as plt  # noqa: E402

TARGET = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE.parent / "target"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else HERE.parent / "golden"
REPO = HERE.parents[1]
SCRUB = REPO / ".claude/skills/reduce-house-fixes/scripts/scrub_producer_metadata.py"

DECKS = ["Civic Center North Deck", "Civic Center South Deck"]
FORECAST_MADE = date(2027, 1, 25)      # the note's date; FES-07 s.3 takes the factor in force that day
BACKTEST_MADE = date(2025, 1, 31)      # the 2025 forecast, made once December 2024 had closed
NOTE_DATE = datetime(2027, 1, 25, 16, 40)
PST = timezone(timedelta(hours=-8))    # the deck sub-meters' clock (nameplate record)
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September",
          "October", "November", "December"]
CONTRACT = [(m, 2027) for m in range(4, 13)] + [(m, 2028) for m in range(1, 4)]   # April 2027 to March 2028


def whole(x):
    return int(math.floor(x + 0.5))


def nearest5(x):
    return int(math.floor(x / 5 + 0.5) * 5)


# ------------------------------------------------------------------------------ the figures
def compute():
    R = rules(str(TARGET))
    P = load(str(TARGET))
    G = Grid(R["holidays"])
    growth = factor_on(R["factors"], FORECAST_MADE)

    rec = of_record(P).copy()
    rec["garage"], rec["position"] = dated_join(P, rec)
    recs = rec[rec["garage"].isin(DECKS) & (rec["plug_in"].str[:4] == "2026")].copy()
    # a charge the 10 a.m. settlement run closed and carried on in a second record is one charge from its first start
    pop = charges(recs)
    # the car each session's permit carries in the contract year: the January 2027 renewal check
    pop["car"] = car_ratings(P, pop, as_of=FORECAST_MADE)
    pop["car_then"] = car_ratings(P, pop)
    draw = np.minimum(R["new_kw"], pop["car"])

    def replay(rate):
        rate = pd.Series(rate, index=pop.index)
        return {g: G.blocks(pop.loc[pop["garage"] == g, "t0"], pop.loc[pop["garage"] == g, "kwh_delivered"],
                            rate[pop["garage"] == g]) for g in DECKS}

    percar = replay(draw)
    tot = percar[DECKS[0]] + percar[DECKS[1]]
    base, t_bind = G.peak(tot)
    j = (t_bind - G.t0) // Q
    monthly = {}
    for m in range(1, 13):
        v, t = G.peak(tot, month=m)
        monthly[m] = {"base": v, "forecast": growth * v, "at": datetime.fromtimestamp(t, LA)}

    rated = replay(np.full(len(pop), R["new_kw"]))
    rung2, t2 = G.peak(rated[DECKS[0]] + rated[DECKS[1]])
    then = replay(np.minimum(R["new_kw"], pop["car_then"]))
    own_date, t4 = G.peak(then[DECKS[0]] + then[DECKS[1]])
    # the note's "the 11 kW cars nearly always are" charged before noon: their morning sessions on the new units
    end = pop["t0"] + pop["kwh_delivered"] / draw * 3600
    noon = pd.to_datetime(pop["plug_in"].str[:10]).map(
        lambda d: datetime(d.year, d.month, d.day, 12, tzinfo=LA).timestamp())
    fast_am = (pop["car"] >= 11.0) & (pop["t0"] < noon)
    fast_done = float((fast_am & (end <= noon)).sum() / fast_am.sum())
    # the 7.2 and 7.7 kW cars' part of the binding quarter-hour (the note's "most of December's peak")
    slow = pop["car"] <= 7.7
    slow_load = sum(G.blocks(pop.loc[slow & (pop["garage"] == g), "t0"], pop.loc[slow & (pop["garage"] == g), "kwh_delivered"],
                             draw[slow & (pop["garage"] == g)])[j] for g in DECKS)

    sp = P["sp"]
    rows = sp[sp["session_id"].isin(set(recs["session_id"]))]
    rows = rows[rows["version"] == rows["session_id"].map(recs.set_index("session_id")["version"])]
    closed = G.readings_load(rows["q"], rows["kwh"])

    units = R["units"]
    div = [f for (a, b), f in R["diversity"].items() if a <= units <= b][0]
    planners = math.ceil(units * R["new_kw"] * div / 5) * 5

    ch = P["checks"].copy()
    ch["checked_on"] = pd.to_datetime(ch["checked_on"]).dt.date
    jan = ch[ch["checked_on"] >= date(2027, 1, 1)].set_index("permit_no")
    before = ch[ch["checked_on"] < date(2027, 1, 1)].sort_values("checked_on").groupby("permit_no").last()
    renewed = sorted(p for p in jan.index if jan.loc[p, "vin"] != before.loc[p, "vin"])
    reg = P["permits"].set_index("permit_no")
    renewal = {
        "deck": sorted(set(reg.loc[renewed, "deck"])),
        "holder": sorted(set(reg.loc[renewed, "holder"])),
        "from": sorted(set(zip(before.loc[renewed, "make"], before.loc[renewed, "model"], before.loc[renewed, "model_year"].astype(int)))),
        "to": sorted(set(zip(jan.loc[renewed, "make"], jan.loc[renewed, "model"], jan.loc[renewed, "model_year"].astype(int)))),
    }
    ref = P["ref"]

    def rating(mk, md, tr, yr):
        hit = ref[(ref["make"].str.upper() == mk) & (ref["model"].str.upper() == md) & (ref["model_year_from"] <= yr)
                  & (ref["model_year_to"] >= yr) & ((ref["trim"] == "") | (ref["trim"].str.upper() == tr))]
        return float(hit["onboard_charger_kw"].iloc[0])
    jan_kw = [rating(mk, md, tr, yr) for mk, md, tr, yr in zip(jan["make"], jan["model"], jan["trim"], jan["model_year"])]
    slow_2027 = sum(1 for x in jan_kw if x <= 7.7)

    F = {
        "growth": growth, "new_kw": R["new_kw"], "units": units, "diversity": div, "planners": planners,
        "base": base, "answer": growth * base, "t_bind": datetime.fromtimestamp(t_bind, LA),
        "north": growth * percar[DECKS[0]][j], "south": growth * percar[DECKS[1]][j],
        "north_base": percar[DECKS[0]][j], "south_base": percar[DECKS[1]][j],
        "monthly": monthly, "rung2": growth * rung2, "rung2_at": datetime.fromtimestamp(t2, LA),
        "day": day_series(G, closed, growth * tot, datetime.fromtimestamp(t_bind, LA).date()),
        "renewed": len(renewed), "slow_2027": slow_2027, "permits_2027": len(jan),
        "own_date": growth * own_date, "own_date_at": datetime.fromtimestamp(t4, LA),
        "renewal": renewal, "slow_share": slow_load / base, "fast_done": fast_done, "kw_2027": sorted(set(jan_kw)),
        "pairs": int((pop["n_rec"] > 1).sum()), "charges_2026": len(pop), "records_2026": len(recs),
    }
    F["filed"] = nearest5(F["answer"])
    F["b3"] = backtest(P, R, G, rec)
    F["b1"] = panel_spans(P, G, rec)[0]
    return F


def day_series(G, closed, forecast, day):
    idx = np.flatnonzero(np.array([d == day for d in G.loc.date]))
    t = [G.loc[i].to_pydatetime() for i in idx]
    return {"t": t, "closed": closed[idx], "forecast": forecast[idx], "day": day}


def backtest(P, R, G, rec):
    """FES-07 applied to 2025 from 2024 at the decks: every dated register assignment, the gateway B sessions at the
    units that reported through it, the fleet card charges the export does not carry, the accepted version of each
    restated session."""
    load_, _, _ = backtest_load(P, G, rec)
    fac = factor_on(R["factors"], BACKTEST_MADE)
    out = []
    for m in range(1, 13):
        b = G.peak(load_, year=2024, month=m)[0]
        a = G.peak(load_, year=2025, month=m)[0]
        fc, ac = whole(fac * b), whole(a)
        out.append({"month": m, "base": b, "factor": fac, "forecast": fc, "recorded": ac,
                    "over": fc - ac, "miss": round(100 * (fc - ac) / ac, 1)})
    return out


def assert_record(F):
    """The build record (DESIGN_NOTE.md, build record; verify_pack.EXPECTED) and the goldens agree."""
    assert abs(F["answer"] - EXPECTED["answer"]) < 1e-6 and F["filed"] == EXPECTED["filed"], F["answer"]
    assert math.ceil(F["answer"] / 5) * 5 == F["filed"]                 # rounding up files the same
    assert F["t_bind"].isoformat() == EXPECTED["binding"], F["t_bind"]
    assert abs(F["north"] - EXPECTED["split"][0]) < 1e-6 and abs(F["south"] - EXPECTED["split"][1]) < 1e-6
    for m in range(1, 13):
        assert abs(F["monthly"][m]["forecast"] - EXPECTED["monthly"][m - 1]) < 1e-6, m
    assert max(F["monthly"], key=lambda m: F["monthly"][m]["forecast"]) == 12
    assert abs(F["own_date"] - EXPECTED["rung4"]) < 1e-6 and F["renewed"] == 18
    # the note's wording about the renewal and the rival peak, back-tested on the record
    rn = F["renewal"]
    assert rn["deck"] == ["North"] and rn["holder"] == ["Larch County Fleet Services"], rn
    assert rn["from"] == [("CHEVROLET", "BOLT EV", 2020)] and rn["to"] == [("CHEVROLET", "BOLT EV", 2023)], rn
    assert F["own_date_at"].month == 2 and F["slow_share"] > 0.5, (F["own_date_at"], F["slow_share"])
    assert F["fast_done"] >= 0.95 and F["kw_2027"] == [7.2, 7.7, 11.0], (F["fast_done"], F["kw_2027"])
    assert abs(F["rung2"] - EXPECTED["rung2"]) < 1e-6 and F["planners"] == EXPECTED["planners"]
    assert F["pairs"] > 2000 and F["records_2026"] - F["charges_2026"] == F["pairs"], (F["pairs"], F["records_2026"])
    assert [b["forecast"] for b in F["b3"]] == EXPECTED["b3_forecast"]
    assert [b["miss"] for b in F["b3"]] == EXPECTED["b3_miss"]
    for panel in ("CP-N", "CP-S"):
        assert [whole(s["gap"]) for s in F["b1"][panel]] == EXPECTED["b1"][panel], panel
    # the chart's marked point is the note's figure
    k = F["day"]["t"].index(F["t_bind"])
    assert abs(F["day"]["forecast"][k] - F["answer"]) < 1e-9
    assert whole(F["north"]) + whole(F["south"]) == whole(F["answer"])


# ------------------------------------------------------------------------------ the chart
INK, INK2, MUTED, GRID, AXIS = "#0b0b0b", "#52514e", "#898781", "#e1e0d9", "#c3c2b7"
BLUE, ORANGE, SURFACE, WASH = "#2a78d6", "#eb6834", "#fcfcfb", "#f0efec"


def chart(F, path):
    plt.rcParams.update({"font.family": ["Inter", "DejaVu Sans"], "font.size": 9.5,
                         "axes.edgecolor": AXIS, "axes.labelcolor": INK2, "xtick.color": INK2, "ytick.color": INK2,
                         "svg.hashsalt": "task117"})
    d = F["day"]
    t = [x.replace(tzinfo=None) for x in d["t"]]
    t_end = t + [t[-1] + timedelta(minutes=15)]
    fig, ax = plt.subplots(figsize=(10, 5.6), dpi=170)
    fig.patch.set_facecolor(SURFACE)
    ax.set_facecolor(SURFACE)
    day0 = datetime.combine(d["day"], datetime.min.time())
    lo, hi = day0 + timedelta(hours=5), day0 + timedelta(hours=21)
    ax.axvspan(day0 + timedelta(hours=12), day0 + timedelta(hours=20), color=WASH, lw=0, zorder=0)
    ax.text(day0 + timedelta(hours=16), 236, "Billing demand counted, 12:00 to 20:00 weekdays",
            ha="center", va="top", color=MUTED, fontsize=8.5)
    ax.stairs(d["closed"], t_end, color=ORANGE, lw=2, baseline=None, zorder=3,
              label="Drawn on the 6.6 kW units, Dec 8, 2026")
    ax.stairs(d["forecast"], t_end, color=BLUE, lw=2, baseline=None, zorder=4,
              label="Forecast on the 11.5 kW units, December 2027 (x 1.12 growth)")
    filed = F["filed"]
    ax.axhline(filed, color=INK, lw=1, zorder=2)
    ax.text(hi - timedelta(minutes=10), filed + 3, f"Contract demand {filed} kW", ha="right", va="bottom",
            color=INK, fontsize=9, fontweight="bold")
    tb = F["t_bind"].replace(tzinfo=None)
    mid = tb + timedelta(minutes=7.5)
    ax.plot([mid], [F["answer"]], "o", ms=8, color=BLUE, mec=SURFACE, mew=2, zorder=5)
    ax.annotate(f"12:00 quarter-hour: {whole(F['answer'])} kW\nsets the contract (North {whole(F['north'])}, "
                f"South {whole(F['south'])})", xy=(mid, F["answer"]), xytext=(tb + timedelta(hours=1, minutes=40), 182),
                color=INK, fontsize=9, ha="left", va="center",
                arrowprops=dict(arrowstyle="-", color=INK2, lw=0.8, shrinkA=2, shrinkB=6))
    ax.set_xlim(lo, hi)
    ax.set_ylim(0, 240)
    ax.xaxis.set_major_locator(mdates.HourLocator(byhour=range(6, 22, 2)))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%H:%M"))
    ax.set_yticks(range(0, 231, 50))
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: f"{v:,.0f}"))
    ax.set_ylabel("Average demand per quarter-hour (kW)")
    ax.grid(axis="y", color=GRID, lw=1)
    ax.set_axisbelow(True)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.spines["left"].set_visible(False)
    ax.tick_params(length=0)
    ax.legend(loc="upper right", bbox_to_anchor=(0.995, 0.36), frameon=False, fontsize=8.8, labelcolor=INK2)
    fig.text(0.055, 0.955, f"File {filed} kW: the decks' busiest billed quarter-hour on the new units comes to "
             f"{whole(F['answer'])} kW at noon", fontsize=13.5, fontweight="bold", color=INK)
    fig.text(0.055, 0.912, "Civic Center North and South Decks combined, by quarter-hour, Tuesday, December 8, 2026, "
             "the basis day for December 2027", fontsize=9.5, color=INK2)
    fig.text(0.055, 0.022, "Source: Curbline settlement and interval exports, 2026 deck charges (records split at the "
             "settlement run rejoined); each charge replayed from its start at the lower of 11.5 kW\nand the onboard "
             "charger rating of the car its permit carries in 2027 (January 2027 vehicle check, Parking Services "
             "reference list), grown by the FES-07 factor of 1.12.\n"
             "Billing demand per NSPL Schedule 26. Energy & Facilities, January 2027.", fontsize=7.6, color=MUTED,
             linespacing=1.4)
    fig.subplots_adjust(left=0.075, right=0.975, top=0.86, bottom=0.19)
    fig.savefig(path, facecolor=SURFACE, metadata={"Software": None})
    plt.close(fig)


# ------------------------------------------------------------------------------ the note
def note(F, path):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    filed, ans = F["filed"], F["answer"]
    north, south = whole(F["north"]), whole(F["south"])
    dec = F["monthly"][12]
    gap = F["planners"] - filed
    rate = 11.40
    body = ParagraphStyle("b", fontName="Helvetica", fontSize=9.6, leading=12.8, spaceAfter=6.5, alignment=TA_LEFT)
    small = ParagraphStyle("s", parent=body, fontSize=7.6, leading=9.6, textColor=colors.HexColor("#3d3d3a"), spaceAfter=2)
    head = ParagraphStyle("h", parent=body, fontName="Helvetica-Bold", fontSize=13.5, leading=16.5, spaceAfter=7)
    lh = ParagraphStyle("lh", parent=body, fontName="Helvetica-Bold", fontSize=10.5, leading=12.5, spaceAfter=0)
    lh2 = ParagraphStyle("lh2", parent=body, fontSize=8.4, leading=10.5, textColor=colors.HexColor("#52514e"), spaceAfter=0)

    def foot(c, doc):
        c.saveState()
        c.setFont("Helvetica", 7.5)
        c.setFillColor(colors.HexColor("#52514e"))
        c.drawString(0.8 * inch, 0.5 * inch, "City of Larch Harbor  |  Energy & Facilities  |  Council packet, "
                                              "February 16, 2027")
        c.drawRightString(7.7 * inch, 0.5 * inch, "Page 1 of 1")
        c.restoreState()

    s = []
    s.append(Paragraph("CITY OF LARCH HARBOR", lh))
    s.append(Paragraph("Public Works, Energy &amp; Facilities Division", lh2))
    s.append(Spacer(1, 9))
    memo = [["To:", "Mayor and City Council"],
            ["From:", "Shelley Tanner, City Energy Manager"],
            ["Date:", "January 25, 2027"],
            ["Re:", "Schedule 1 contract demand, NSPL service agreement for the Civic Center decks"]]
    mt = Table(memo, colWidths=[0.6 * inch, 6.2 * inch])
    mt.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 9.2), ("FONT", (0, 0), (0, -1), "Helvetica-Bold", 9.2),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 1.2), ("TOPPADDING", (0, 0), (-1, -1), 1.2),
                            ("LEFTPADDING", (0, 0), (-1, -1), 0),
                            ("LINEBELOW", (0, -1), (-1, -1), 0.6, colors.HexColor("#898781"))]))
    s.append(mt)
    s.append(Spacer(1, 10))
    s.append(Paragraph(f"We should file {filed} kW as the contract demand for the first contract year", head))
    s.append(Paragraph(
        f"I recommend Council approve the agreement with <b>{filed} kW</b> entered in Schedule 1 as the contract demand "
        f"for April 2027 through March 2028. That is the highest billing demand I forecast for the new service in any "
        f"month of the contract year: <b>{whole(dec['forecast'])} kW in December 2027</b>, set in the quarter-hour "
        f"beginning at noon. In that quarter-hour North Deck carries {north} kW and South Deck {south} kW. Schedule 1 "
        f"takes whole multiples of 5 kW, and {dec['forecast']:.1f} kW comes to {filed} kW whether it is rounded to the "
        f"nearest step or up.", body))
    s.append(Paragraph(
        "The forecast follows FES-07: each contract month comes from the same month of 2026 at the equipment the "
        "service will supply, grown by the 2027 factor of 1.12. Rather than scale up the old units' peaks, I replayed "
        "every 2026 charge at the decks on the Exhibit A units, each car taking its 2026 energy from the start of its "
        "charge at the lower of 11.5 kW and its own onboard charger rating. Curbline's 10 a.m. settlement run splits a "
        f"charge still running into two session records; I joined the {F['pairs']:,} pairs back into one charge first, "
        "because record by record every car still charging at the run would restart at full power. At the old units' "
        "6.6 kW the replay reproduces every 2026 interval reading Curbline settled at the decks.", body))
    s.append(Paragraph(
        f"The car that counts is the one each permit carries in the contract year: in January Larch County Fleet "
        f"Services moved {F['renewed']} North Deck permits from 2020 Bolt EVs (7.2 kW) to 2023 Bolt EVs (11 kW). With "
        f"the 2026 cars the year would peak at {whole(F['own_date'])} kW in February. Ricardo Moore expects everyone "
        f"charged before lunch, and the 11 kW cars nearly always are, but {F['slow_2027']} of the {F['permits_2027']} "
        f"permit vehicles are listed at 7.2 or 7.7 kW; the ones arriving mid-morning still charge at noon, and most of "
        f"December's peak is theirs.", body))
    gap = F["planners"] - filed
    s.append(Paragraph(
        f"North Sound Power &amp; Light's planners size an EV service at nameplate times the diversity factor in "
        f"Section 7 of their planning guide: {F['units']} units at 11.5 kW times {F['diversity']:.2f}, to the next "
        f"5 kW above, is <b>{F['planners']} kW</b>, which I expect Paul Henderson to propose at the service review. "
        f"Ours is <b>{gap} kW lower</b>, worth ${gap * rate * 12:,.0f} a year at Schedule 26's ${rate:.2f} a month per "
        f"contracted kW, and Section 4 of the agreement makes the contract demand the maximum billing demand we expect. "
        f"A month above {filed} kW would reset the contract to that month's demand, rounded up to the next 5 kW, for "
        f"twelve months.", body))
    rows = [["Contract month", "Built from", "Forecast billing demand (kW)"]]
    for m, y in CONTRACT:
        v = F["monthly"][m]
        rows.append([f"{MONTHS[m - 1]} {y}", f"{MONTHS[m - 1][:3]} 2026", f"{whole(v['forecast'])}"])
    rows.append(["Contract demand, first contract year", "", f"{filed}"])
    tb = Table(rows, colWidths=[2.6 * inch, 1.3 * inch, 2.0 * inch], hAlign="LEFT")
    dec_row = 1 + CONTRACT.index((12, 2027))
    tb.setStyle(TableStyle([
        ("FONT", (0, 0), (-1, -1), "Helvetica", 8.6), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.6),
        ("FONT", (0, -1), (-1, -1), "Helvetica-Bold", 8.6), ("FONT", (0, dec_row), (-1, dec_row), "Helvetica-Bold", 8.6),
        ("ALIGN", (2, 0), (2, -1), "RIGHT"), ("TOPPADDING", (0, 0), (-1, -1), 1.1), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.1),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#52514e")),
        ("LINEABOVE", (0, -1), (-1, -1), 0.6, colors.HexColor("#52514e")),
        ("BACKGROUND", (0, dec_row), (-1, dec_row), colors.HexColor("#eef3fb"))]))
    s.append(tb)
    s.append(Spacer(1, 3))
    s.append(Paragraph("Source: Curbline settlement and interval exports, sessions plugged in January to December 2026 "
                       "at the 32 deck units, deliveries through January 18, 2027; Parking Services permit registry, "
                       "vehicle checks (renewals through January 15, 2027) and vehicle reference list.", small))
    s.append(Spacer(1, 5))
    s.append(Paragraph("Billing demand is as Schedule 26, Section 3 defines it: the highest average kW in any "
                       "fifteen-minute interval beginning 12:00 noon through 7:45 p.m., Monday to Friday, outside the "
                       "listed holidays. The new service has one meter, so the decks' loads are summed before the "
                       "highest quarter-hour is taken. Monthly figures are rounded to whole kW; the factor of 1.12 is "
                       "FES-07 Table 1, adopted September 15, 2026.", small))
    doc = SimpleDocTemplate(str(path), pagesize=letter, leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                            topMargin=0.65 * inch, bottomMargin=0.75 * inch, title="Contract demand, Civic Center decks",
                            author="Shelley Tanner", subject="Schedule 1 contract demand", creator="City of Larch Harbor",
                            invariant=1)
    doc.build(s, onFirstPage=foot)


# ------------------------------------------------------------------------------ the workbook
def workbook(F, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    hdr_font = Font(name="Calibri", bold=True, size=10)
    hdr_fill = PatternFill("solid", fgColor="DCE6F1")
    thin = Side(style="thin", color="9A9A9A")
    base_font = Font(name="Calibri", size=10)
    title_font = Font(name="Calibri", bold=True, size=12)
    note_font = Font(name="Calibri", italic=True, size=9, color="555555")

    def sheet(ws, title, subtitle, headers, rows, widths, fmts, note_lines=()):
        ws["A1"], ws["A1"].font = title, title_font
        ws["A2"], ws["A2"].font = subtitle, note_font
        r0 = 4
        for c, h in enumerate(headers, 1):
            cell = ws.cell(r0, c, h)
            cell.font, cell.fill = hdr_font, hdr_fill
            cell.alignment = Alignment(wrap_text=True, vertical="bottom", horizontal="left" if c == 1 else "right")
            cell.border = Border(bottom=thin)
        for i, row in enumerate(rows, r0 + 1):
            for c, v in enumerate(row, 1):
                cell = ws.cell(i, c, v)
                cell.font = base_font
                if fmts[c - 1]:
                    cell.number_format = fmts[c - 1]
                cell.alignment = Alignment(horizontal="left" if isinstance(v, str) else "right")
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.row_dimensions[r0].height = 42
        ws.freeze_panes = ws.cell(r0 + 1, 2)
        last = r0 + len(rows)
        for k, line in enumerate(note_lines):
            ws.cell(last + 2 + k, 1, line).font = note_font
        ws.print_title_rows = f"{r0}:{r0}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToHeight = 0
        ws.print_area = f"A1:{get_column_letter(len(headers))}{last + 1 + len(note_lines)}"
        return last

    # 1. the 2025 record of the method
    ws = wb.active
    ws.title = "2025 back-test"
    rows = [[f"{MONTHS[b['month'] - 1]} 2025", f"{MONTHS[b['month'] - 1][:3]} 2024", round(b["base"], 1), b["factor"],
             b["forecast"], b["recorded"], b["over"], round(b["miss"] / 100, 4)] for b in F["b3"]]
    sheet(ws, "FES-07 applied to 2025: forecast from 2024 against recorded billing demand, Civic Center decks",
          "Billing demand per NSPL Schedule 26, Section 3, both decks summed per quarter-hour. Prepared January 2027.",
          ["Month", "Base month", "Base billing demand (kW)", "Factor", "Forecast (kW)", "Recorded billing demand (kW)",
           "Over (+) / under (-) (kW)", "Error, % of recorded"],
          rows, [15, 11, 13, 8, 10, 14, 13, 12], [None, None, "0.0", "0.00", "0", "0", '+0;-0;0', '+0.0%;-0.0%;0.0%'],
          ["Factor 1.08 is the FES-07 Table 1 factor in force when the 2025 forecast was made from the closed 2024 "
           "months (adopted September 10, 2024); the 1.09 revision of April 8, 2025 came later.",
           "Forecast and recorded demand in whole kW and error as forecast less recorded, as a percentage of recorded, "
           "per FES-07 Section 4.",
           "2024 base includes the January to April sessions at the units that reported through gateway B, and each "
           "station identifier is placed by its dated register assignment (identifiers changed April 1, 2025).",
           "Fleet card charges at the decks before the April 1, 2025 platform move settled through the fleet card "
           "processor and are not in the Curbline export; they are added from the fleet card transactions, matched to "
           "the export on NETWORK_REF so no charge counts twice. The processor stamps START and END in UTC (on every "
           "charge both files carry, START is the export's plug-in to the second), so they are placed in local time.",
           "Restated sessions: the version Parking Services marked ACCEPTED, otherwise the version delivered first; "
           "re-delivered records counted once per authorization code."])

    # 2. the 2026 basis against the panel meters
    ws = wb.create_sheet("2026 panel check")
    rows = []
    for panel, deck in (("CP-N", "North"), ("CP-S", "South")):
        for s in F["b1"][panel]:
            rows.append([f"{panel} ({deck})", s["meter"], s["read"], s["date"], s["time"],
                         s["start"].strftime("%b %d %H:%M ") + s["start"].tzname(),
                         s["end"].strftime("%b %d %H:%M ") + s["end"].tzname(), round(s["metered"], 1),
                         round(s["sessions"], 1), whole(s["gap"])])
    last = sheet(ws, "Deck charging panels, 2026: metered energy the settled sessions do not account for",
                 "Each 2026 read of the two sub-meters against the Curbline interval readings on the units fed from "
                 "that panel between consecutive reads.",
                 ["Panel", "Meter", "Read", "Read date", "Read time (meter clock)", "Span from (local time)",
                  "Span to (local time)", "Metered kWh", "Session kWh on panel", "Unaccounted kWh"],
                 rows, [13, 9, 6, 11, 12, 17, 17, 12, 13, 12],
                 [None, None, "0", "yyyy-mm-dd", None, None, None, "#,##0.0", "#,##0.0", "#,##0"],
                 ["The sub-meters keep Pacific Standard Time all year (DST adjustment disabled, nameplate record), so "
                  "read times from March to October are an hour behind local time.",
                  "Session kWh: Curbline readings for every quarter-hour inside the span; in the quarter-hour a read "
                  "falls inside, each session's constant 6.6 kW draw up to the read time. Weekend and Schedule 26 "
                  "holiday charging is free to permit holders and is not settled, so it comes from Curbline's "
                  "courtesy-session report (none of those sessions spans a read).",
                  "On December 31 the South panel was read twice; the later read (09:52) stands, per the log.",
                  "Unit N-11 was on the North Deck house panel from June 1 to July 12, 2026 (WO-26-0418, panel "
                  "schedule), so its sessions in that period are not on CP-N.",
                  "The unaccounted energy is mostly each panel's roof-level pole lighting (circuit 33)."])
    for r in range(5, last + 1):
        ws.cell(r, 4).alignment = Alignment(horizontal="right")
    ws.auto_filter.ref = f"A4:J{last}"

    # 3. the contract year, as filed
    ws = wb.create_sheet("Contract year")
    rows = []
    for m, y in CONTRACT:
        v = F["monthly"][m]
        rows.append([f"{MONTHS[m - 1]} {y}", f"{MONTHS[m - 1][:3]} 2026", v["at"].strftime("%a %b %d, %H:%M"),
                     round(v["base"], 1), F["growth"], whole(v["forecast"])])
    last = sheet(ws, f"Civic Center decks service, first contract year: contract demand {F['filed']} kW",
                 "Forecast billing demand per contract month on the Exhibit A units, from the same month of 2026 "
                 "(FES-07 Section 2).",
                 ["Contract month", "Built from", "Highest billed quarter-hour (start)", "2026 replay (kW)", "Factor",
                  "Forecast billing demand (kW)"],
                 rows, [17, 11, 20, 12, 8, 14], [None, None, None, "0.0", "0.00", "0"],
                 ["Replay: each 2026 deck charge from its start until its delivered kWh, at the lower of 11.5 kW and "
                  "the onboard charger rating of the car its permit carries in 2027 (January 2027 vehicle check, "
                  "vehicle reference list). A charge the 10 a.m. settlement run split into two session records is "
                  "replayed as one.",
                  f"Contract demand: the highest month ({whole(F['monthly'][12]['forecast'])} kW, December 2027) in "
                  f"whole multiples of 5 kW: {F['filed']} kW."])
    dec_row = 5 + CONTRACT.index((12, 2027))
    for c in range(1, 7):
        ws.cell(dec_row, c).font = Font(name="Calibri", size=10, bold=True)

    # 4. notes
    ws = wb.create_sheet("Notes")
    notes = [("Prepared by", "Shelley Tanner, City Energy Manager, Energy & Facilities Division"),
             ("Prepared", "January 25, 2027, for the Council packet of February 16, 2027"),
             ("Billing demand", "Highest average kW in any 15-minute interval beginning 12:00 noon to 7:45 p.m., "
                                "Monday to Friday, excluding the Schedule 26 holidays (NSPL Schedule 26, Section 3)."),
             ("Demand from readings", "Four times the quarter-hour kWh (Curbline field notes)."),
             ("Sessions", "Curbline settlement export, all deliveries through January 18, 2027; interval export for "
                          "the same records. The 10 a.m. settlement run closes a session still charging and the charge "
                          "carries on in a new record from that second at the same unit, under the same permit or fleet card; "
                          "the forecast replays charges, so those pairs are joined back into one."),
             ("Courtesy sessions", "curbline_courtesy_sessions_civic_decks_2024-2026.csv, for the panel check (weekend "
                                   "and holiday charging)."),
             ("Garage and unit", "station_register.csv, by the assignment in service on the session date."),
             ("Cars", "permit_vehicle_checks.csv, the check at the January 2027 renewal for the car each permit "
                      "carries in the contract year; onboard charger rating from vehicle_reference_list.csv; fleet "
                      "cards through city_fleet_roster.csv."),
             ("Fleet card charges", "fleet_card_ev_transactions_2024-2026.csv for the charges the Curbline export "
                                    "does not carry (before April 1, 2025)."),
             ("Panel circuits", "deck_panel_circuit_schedule.csv, by the assignment in force on the reading date."),
             ("Meter reads", "deck_panel_meter_log_2024-2026.xlsx; clock per deck_submeter_nameplates.csv."),
             ("Growth", "FES-07 Rev. 4, Table 1: 1.12 for forecasts made from September 15, 2026; 1.08 for the 2025 "
                        "forecast."),
             ("Rounding", "kW and kWh to whole numbers in the filed figures; errors to one decimal of a percent.")]
    ws["A1"], ws["A1"].font = "Notes and sources", title_font
    for i, (k, v) in enumerate(notes, 3):
        ws.cell(i, 1, k).font = Font(name="Calibri", bold=True, size=10)
        c = ws.cell(i, 2, v)
        c.font, c.alignment = base_font, Alignment(wrap_text=True, vertical="top")
        ws.cell(i, 1).alignment = Alignment(vertical="top")
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 96

    wb.properties.creator = "Shelley Tanner"
    wb.properties.lastModifiedBy = "Shelley Tanner"
    wb.properties.title = "Civic Center decks service demand"
    wb.properties.created = NOTE_DATE
    wb.properties.modified = NOTE_DATE
    wb.active = 0
    wb.save(path)


def repack(path, when):
    """Fixed entry times so two runs are byte-identical (H7)."""
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        data = {i.filename: z.read(i.filename) for i in infos}
    # openpyxl stamps dcterms:modified with the build clock on save; the workbook is the note's
    iso = when.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    data["docProps/core.xml"] = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rb"\g<1>" + iso,
                                       data["docProps/core.xml"])
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for i in infos:
            zi = zipfile.ZipInfo(i.filename, date_time=when.timetuple()[:6])
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            z.writestr(zi, data[i.filename])


# ------------------------------------------------------------------------------ main
def main():
    F = compute()
    assert_record(F)
    OUT.mkdir(parents=True, exist_ok=True)
    chart(F, OUT / "deck_load_day.png")
    note(F, OUT / "contract_demand_note.pdf")
    workbook(F, OUT / "civic_service_demand.xlsx")
    repack(OUT / "civic_service_demand.xlsx", NOTE_DATE)
    a = subprocess.run([sys.executable, str(SCRUB), str(OUT), "--apply", "--producer", "City of Larch Harbor",
                        "--stamp", "2027-01-25", "--floor", "2026-01-01", "--ceiling", "2027-01-25"],
                       capture_output=True, text=True)
    assert a.returncode == 0, a.stdout + a.stderr
    repack(OUT / "civic_service_demand.xlsx", NOTE_DATE)
    audit = subprocess.run([sys.executable, str(SCRUB), str(OUT), "--floor", "2026-01-01", "--ceiling", "2027-01-25"],
                           capture_output=True, text=True)
    assert audit.returncode == 0 and "clean" in audit.stdout, audit.stdout + audit.stderr
    stamp = NOTE_DATE.timestamp()
    for f in os.listdir(OUT):
        os.utime(OUT / f, (stamp, stamp))

    print(f"CALL  contract demand {F['filed']} kW (unrounded {F['answer']:.3f}; rounded up "
          f"{math.ceil(F['answer'] / 5) * 5})")
    print("\ncritical components")
    print(f"  draw on the new units      min({F['new_kw']} kW, onboard rating of the 2027 permit car); "
          f"{F['renewed']} permits renewed onto 11 kW cars; 2027 permit vehicles at 7.2 or 7.7 kW: "
          f"{F['slow_2027']} of {F['permits_2027']}")
    print(f"  cars on their 2026 dates   {F['own_date']:.3f} kW, set {F['own_date_at']:%d %b %Y %H:%M} (rival)")
    print(f"  2026 billed-hours maximum  {F['base']:.1f} kW at {F['t_bind']:%H:%M} on {F['t_bind']:%a %d %b %Y} "
          f"(North {F['north_base']:.1f}, South {F['south_base']:.1f})")
    print(f"  growth factor              {F['growth']:.2f}")
    print(f"  December 2027              {F['monthly'][12]['forecast']:.1f} kW, the highest contract month")
    print(f"\ndeck split after growth     North {F['north']:.3f} ({whole(F['north'])}), South {F['south']:.3f} "
          f"({whole(F['south'])})")
    print(f"NSPL planners' sizing       {F['units']} x {F['new_kw']} x {F['diversity']:.2f} = "
          f"{F['units'] * F['new_kw'] * F['diversity']:.1f} -> {F['planners']} kW; gap {F['planners'] - F['filed']} kW")
    print(f"replay at the rating (rival) {F['rung2']:.2f} kW, set {F['rung2_at']:%d %b %Y %H:%M}")
    print("\ncontract year (whole kW)")
    for m, y in CONTRACT:
        v = F["monthly"][m]
        print(f"  {MONTHS[m - 1][:3]} {y}  {whole(v['forecast']):>4}   ({v['forecast']:.3f}, from {v['at']:%d %b %H:%M})")
    print("\n2025 back-test (forecast kW, recorded kW, error %)")
    for b in F["b3"]:
        print(f"  {MONTHS[b['month'] - 1][:3]} 2025  {b['forecast']:>4}  {b['recorded']:>4}  {b['miss']:+.1f}")
    print("\n2026 panel check (unaccounted kWh)")
    for panel in ("CP-N", "CP-S"):
        print(f"  {panel}  " + "  ".join(str(whole(s['gap'])) for s in F["b1"][panel]))


if __name__ == "__main__":
    main()
