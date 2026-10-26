"""The independent verifier: reads only the bundle under target/ (csv, xlsx, pdf, docx), shares no code or state with
the generator or golden.py, recomputes the planning figure, its parts and every ask figure on its own path (the
standard library's csv and datetime, no pandas), and compares them with generator/figures.json.

    python3 verify.py [target_dir]
"""
import collections
import csv
import datetime as dt
import json
import os
import re
import sys
import zipfile
from zoneinfo import ZoneInfo

import openpyxl
import pypdf
import docx

HERE = os.path.dirname(os.path.abspath(__file__))
UK = ZoneInfo("Europe/London")
UTC = dt.timezone.utc
FAILS = []


def ok(cond, label):
    print(("ok    " if cond else "FAIL  ") + label)
    if not cond:
        FAILS.append(label)


def rows(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def utc(s):
    return dt.datetime.strptime(s, "%Y-%m-%dT%H:%M:00Z").replace(tzinfo=UTC)


def uk(s):
    return dt.datetime.strptime(s, "%d/%m/%Y %H:%M").replace(tzinfo=UK).astimezone(UTC)


def fy(t):
    t = t.astimezone(UK)
    return t.year if t.month >= 4 else t.year - 1


def pdf_text(path):
    return "\n".join(p.extract_text() or "" for p in pypdf.PdfReader(path).pages)


def docx_text(path):
    return "\n".join(p.text for p in docx.Document(path).paragraphs)


def main(tgt):
    fig = json.load(open(os.path.join(HERE, "figures.json")))
    names = sorted(os.listdir(tgt))
    pick = lambda pre: os.path.join(tgt, [n for n in names if n.startswith(pre)][0])
    # documents read back, each with a positive control on the same extractor
    brief = docx_text(pick("tanner_street_hub_planning_brief"))
    ok("opens on 1 April 2027" in brief and len(brief) > 1500, "brief read back (positive control)")
    model = pdf_text(pick("custody_demand_model"))
    ok("capped at 10 per cent a year" in model.replace("\n", " "), "model note read back (positive control)")
    wb = openpyxl.load_workbook(pick("custody_suite_allocation"), read_only=True)
    alloc = [r for r in wb["Allocation"].iter_rows(values_only=True)]
    notes = " ".join(str(c) for r in wb["Notes"].iter_rows(values_only=True) for c in r if c)
    ok("whichever force or agency" in notes, "allocation notes read back (positive control)")
    hdr = alloc[0]
    hub = sorted(r[hdr.index("ward")] for r in alloc[1:] if r[hdr.index("suite_from_1_april_2027")] == "Tanner Street")
    home = {r[hdr.index("ward")]: r[hdr.index("suite_to_31_march_2027")] for r in alloc[1:]}
    ok(len(hub) == 6, "six wards move to Tanner Street")

    # the force's arrests in the six wards by UK financial year
    arr = rows(pick("calderwick_police_arrests"))
    per = collections.Counter((r["ward"], fy(utc(r["arrested_utc"]))) for r in arr if r["ward"] in hub)
    own = 0.0
    for w in hub:
        ch = (per[(w, 2025)] / per[(w, 2022)]) ** (1 / 3) - 1
        own += per[(w, 2025)] * (1 + max(-0.1, min(0.1, ch))) ** 2
    ok(abs(own - fig["components.own_arrests_model"]) < 0.05, "model on the six wards %.1f" % own)
    ok(sum(per[(w, 2025)] for w in hub) == fig["components.own_arrests_2025_26"], "six wards' 2025-26 arrests")
    for y in range(2021, 2026):
        ok(sum(per[(w, y)] for w in hub) == fig["chart.six_wards_arrests.%d-%02d" % (y, (y + 1) % 100)], "chart %d" % y)
    # the model note's validation table reproduces from the arrest records
    by_suite = collections.defaultdict(lambda: collections.Counter())
    for r in arr:
        by_suite[home[r["ward"]]][(r["ward"], fy(utc(r["arrested_utc"])))] += 1
    flat = re.sub(r"\s+", " ", model)
    for s in ("Northgate", "Eastmoor", "Southfields", "Wendle Cross"):
        c = by_suite[s]
        wards = {w for w, _ in c}
        fc = sum(c[(w, 2024)] * (1 + max(-0.1, min(0.1, (c[(w, 2024)] / c[(w, 2021)]) ** (1 / 3) - 1))) for w in wards)
        act = sum(c[(w, 2025)] for w in wards)
        ok(("%s %s %s %.3f" % (s, "{:,}".format(round(fc)), "{:,}".format(act), fc / act)) in flat,
           "validation row %s reproduces" % s)

    # British Transport Police's crime at the hub's stations, at the force's own arrests per crime of each type
    lsoa = {r["lsoa21_code"]: r["ward"] for r in rows(pick("lsoa21_to_ward_lookup"))}
    pol = rows(pick("police_uk_street_crime"))
    hp = collections.Counter(r["Crime type"] for r in pol if r["Reported by"] == "Calderwick Police")
    bt = collections.Counter(r["Crime type"] for r in pol if r["Reported by"] == "British Transport Police"
                             and lsoa[r["LSOA code"]] in hub)
    ok(sum(bt.values()) == fig["components.btp_crimes_hub_2025_26"], "BTP crimes in the six wards %d" % sum(bt.values()))
    a25 = collections.Counter(r["offence_group"] for r in arr if fy(utc(r["arrested_utc"])) == 2025)
    warr = a25["Warrant, recall or breach"] / (sum(a25.values()) - a25["Warrant, recall or breach"])
    btp = sum(n * a25[t] / hp[t] for t, n in bt.items()) * (1 + warr)
    ok(abs(btp - fig["components.btp_detentions"]) < 0.05, "BTP detentions %.1f" % btp)
    tot = own + btp
    ok(int(tot / 250 + 0.5) * 250 == fig["call.planning_figure"], "planning figure %d (%.1f)" % (int(tot / 250 + 0.5) * 250, tot))
    ok(int(own / 250 + 0.5) * 250 == fig["components.default_answer"], "the model alone prints %d" % (int(own / 250 + 0.5) * 250))

    # ask A: remand holds by suite, defendants charged in 2025-26
    rem = {r["arrest_ref"]: r for r in arr if r["outcome"] == "Charged - remanded in custody"}
    for r in rem.values():
        h, m = map(int, r["detention_hhmm"].split(":"))
        r["_book"] = utc(r["booked_in_utc"])
        r["_charge"] = r["_book"] + dt.timedelta(minutes=60 * h + m)
    by_ch = collections.defaultdict(list)
    for ref, r in rem.items():
        by_ch[(r["custody_suite"], r["_charge"])].append(ref)
    files = rows(pick("cju_remand_files"))
    own_det = {}
    for f in files:
        L = rem.get(f["linked_arrest"])
        if L is not None and L["_charge"] == uk(f["charged"]) and L["custody_suite"] == f["custody_suite"]:
            own_det[f["urn"]] = f["linked_arrest"]
    claimed = set(own_det.values())
    for f in files:
        if f["urn"] not in own_det:
            hit = [ref for ref in by_ch[(f["custody_suite"], uk(f["charged"]))] if ref not in claimed]
            if len(hit) != 1:
                ok(False, "file %s has one detention (%s)" % (f["urn"], hit))
                continue
            own_det[f["urn"]] = hit[0]
    ok(len(set(own_det.values())) == len(files), "every file has its own detention")
    left = {ref for ref in rem if ref not in set(own_det.values())}
    inc = collections.defaultdict(list)
    for ref in left:
        r = rem[ref]
        inc[(r["arrested_utc"], r["ward"], r["offence_group"])].append(ref)
    hold = collections.defaultdict(float)
    over = collections.Counter()
    for f in files:
        L = rem[own_det[f["urn"]]]
        hand = uk(f["handed_to_escort"])
        for ref in [own_det[f["urn"]]] + inc.get((L["arrested_utc"], L["ward"], L["offence_group"]), []):
            r = rem[ref]
            if fy(r["_charge"]) != 2025:
                continue
            hold[f["custody_suite"]] += (hand - r["_charge"]).total_seconds() / 3600
            over[f["custody_suite"]] += (hand - r["_book"]).total_seconds() > 24 * 3600
    for s in ("Northgate", "Eastmoor", "Southfields", "Wendle Cross"):
        ok(round(hold[s]) == fig["ask_a.%s.cell_hours" % s], "ask A %s cell hours %d" % (s, round(hold[s])))
        ok(over[s] == fig["ask_a.%s.over_24_hours" % s], "ask A %s over 24 hours %d" % (s, over[s]))

    # ask B: initial health assessments by the suite they took place in, begun by day and at night
    bk = collections.defaultdict(list)
    for r in arr:
        bk[r["custody_suite"]].append(utc(r["booked_in_utc"]))
    shut = {}
    for s, ts in bk.items():
        ts.sort()
        for t0, t1 in zip(ts, ts[1:]):
            if t1 - t0 > dt.timedelta(days=3):
                during = collections.Counter(r["custody_suite"] for r in arr if home[r["ward"]] == s
                                             and t0 < utc(r["booked_in_utc"]) < t1)
                shut[s] = (t0.astimezone(UK).replace(tzinfo=None), t1.astimezone(UK).replace(tzinfo=None),
                           during.most_common(1)[0][0])
    ok(sorted(shut) == ["Eastmoor", "Wendle Cross"], "two suites closed in the records: %s" % sorted(shut))
    hc = rows(pick("custody_healthcare_contacts"))
    ep = collections.defaultdict(list)
    for r in hc:
        if r["contact_type"] == "Initial health assessment":
            ep[r["episode_ref"]].append(r)
    grid = collections.Counter()
    for lines in ep.values():
        lines.sort(key=lambda r: r["started"])
        prev = None
        for r in lines:
            s0 = dt.datetime.strptime(r["started"], "%Y-%m-%d %H:%M")
            cont = prev is not None and prev == r["started"] and s0.minute == 0 and s0.hour in (7, 19)
            prev = r["ended"]
            if cont:
                continue
            place = r["suite"]
            if place in shut and shut[place][0] <= s0 < shut[place][1]:
                place = shut[place][2]
            grid[(place, "day" if 7 <= s0.hour < 19 else "night")] += 1
    for s in ("Northgate", "Eastmoor", "Southfields", "Wendle Cross"):
        for k in ("day", "night"):
            ok(grid[(s, k)] == fig["ask_b.%s.%s" % (s, k)], "ask B %s %s %d" % (s, k, grid[(s, k)]))
    print("%d checks failed" % len(FAILS) if FAILS else "verify: every figure reproduces from the bundle")
    return 1 if FAILS else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(HERE), "target")))
