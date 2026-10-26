"""The golden deliverables, from golden.figures(): tanner_street_planning_figure_2027-28.xlsx (the planning figure and
how it is built, the remand holds, the initial health assessments) and hub_planning_note.docx (the note with its
chart). Every figure is carried from figures(), never typed."""
import io
import os
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt, RGBColor
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt                                         # noqa: E402
import matplotlib.ticker                                                # noqa: E402
import openpyxl                                                         # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import golden as GD        # noqa: E402
import postprocess as PP   # noqa: E402

F_XLSX = "tanner_street_planning_figure_2027-28.xlsx"
F_DOCX = "hub_planning_note.docx"
PLANNER, PROGRAMME = "Emma Roberts", "Liam Clayton"
TYPES_ORDER = ["Violence and sexual offences", "Public order", "Drugs", "Possession of weapons", "Robbery", "Burglary",
               "Criminal damage and arson", "Shoplifting", "Theft from the person", "Other theft", "Vehicle crime",
               "Bicycle theft", "Other crime"]
BLUE, ORANGE, INK, GREY = "#2a78d6", "#eb6834", "#1f2428", "#5b6670"


def n0(x):
    return "{:,}".format(int(round(x)))


def ward_rows(M):
    rows = []
    for w in sorted(M["by_ward_fy"]):
        y = M["by_ward_fy"][w]
        ch = (y[2025] / y[2022]) ** (1 / 3) - 1
        chc = max(-0.10, min(0.10, ch))
        rows.append((w, y[2022], y[2025], chc, y[2025] * (1 + chc) ** 2))
    return rows


def btp_rows(M):
    out = []
    for t in TYPES_ORDER:
        n = M["b25"].get(t, 0)
        r = M["rate_t"].get(t, 0.0)
        out.append((t, n, r, n * r))
    return out


# ------------------------------------------------------------------------------------------------------------------
def write_xlsx(R, path):
    M, A, B = R["M"], R["A"], R["B"]
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1B365D")
    sub = PatternFill("solid", fgColor="E8EDF3")
    thin = Side(style="thin", color="A7B0B8")
    top = Border(top=thin)

    def header(ws, row, labels):
        for j, t in enumerate(labels, 1):
            c = ws.cell(row=row, column=j, value=t or None)
            c.font, c.fill = head, fill
            c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left" if j == 1 else "right")

    ws = wb.active
    ws.title = "Planning figure"
    ws["A1"] = "Tanner Street custody hub, 2027-28: a planning figure of %s detentions" % n0(M["golden"])
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Detentions the hub will book from 1 April 2027 to 31 March 2028. Custody Services Planning, 26 October 2026."
    ws["A4"], ws["B4"] = "Planning figure for 2027-28 (nearest 250)", M["golden"]
    ws["A4"].font = ws["B4"].font = Font(bold=True, size=12)
    ws["B4"].number_format = "#,##0"
    ws["A5"], ws["B5"] = "Unrounded", round(M["golden_raw"], 1)
    ws["B5"].number_format = "#,##0.0"

    r = 7
    ws.cell(row=r, column=1, value="1. Arrests by Calderwick Police officers in the six wards, custody demand model v2.0").font = bold
    r += 1
    header(ws, r, ["Ward", "Arrests 2022-23", "Arrests 2025-26", "Annual change (capped at 10%)", "2027-28"])
    for w, a22, a25, ch, f27 in ward_rows(M):
        r += 1
        ws.cell(row=r, column=1, value=w)
        ws.cell(row=r, column=2, value=a22).number_format = "#,##0"
        ws.cell(row=r, column=3, value=a25).number_format = "#,##0"
        ws.cell(row=r, column=4, value=round(ch, 4)).number_format = "0.0%"
        ws.cell(row=r, column=5, value=round(f27, 1)).number_format = "#,##0.0"
    r += 1
    for j, v in ((1, "Six wards"), (2, sum(x[1] for x in ward_rows(M))), (3, M["own_2025"]), (5, round(M["F"]["F1 the model as filed"], 1))):
        c = ws.cell(row=r, column=j, value=v)
        c.font, c.border = bold, top
        if j > 1:
            c.number_format = "#,##0.0" if j == 5 else "#,##0"
    r_own = r

    r += 2
    ws.cell(row=r, column=1, value="2. Arrests by British Transport Police officers at Calderwick and Kingsway stations").font = bold
    r += 1
    ws.cell(row=r, column=1, value="BTP has no cells of its own; its officers book the people they arrest into the "
                                   "suite serving the ward (suite allocation, note 2). Its arrests are in none of the "
                                   "force's records, so they are sized from the crimes BTP recorded at the two stations "
                                   "in 2025-26 at the force's own 2025-26 arrests per recorded crime of each type.")
    ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.row_dimensions[r].height = 44
    r += 1
    header(ws, r, ["Crime type", "BTP crimes 2025-26", "Force arrests per crime 2025-26", "", "BTP arrests"])
    for t, n, rt, a in btp_rows(M):
        r += 1
        ws.cell(row=r, column=1, value=t)
        ws.cell(row=r, column=2, value=n).number_format = "#,##0"
        ws.cell(row=r, column=3, value=round(rt, 4)).number_format = "0.000"
        ws.cell(row=r, column=5, value=round(a, 1)).number_format = "#,##0.0"
    r += 1
    per_type = M["B"]["B1 per-type rates"]
    ws.cell(row=r, column=1, value="Arrests for recorded crime").font = bold
    ws.cell(row=r, column=2, value=M["btp_hub_crimes"]).number_format = "#,##0"
    ws.cell(row=r, column=5, value=round(per_type, 1)).number_format = "#,##0.0"
    for j in (1, 2, 5):
        ws.cell(row=r, column=j).border = top
    r += 1
    ws.cell(row=r, column=1, value="Warrant, recall or breach (the force's %.1f%% of arrests for crime)" % (100 * M["warrant_share"]))
    ws.cell(row=r, column=5, value=round(per_type * M["warrant_share"], 1)).number_format = "#,##0.0"
    r += 1
    ws.cell(row=r, column=1, value="BTP arrests at the two stations").font = bold
    c = ws.cell(row=r, column=5, value=round(M["B"]["B2 per-type rates with the warrant share"], 1))
    c.number_format, c.font = "#,##0.0", bold
    r += 1
    st = M["btp_hub_by_station"]
    ws.cell(row=r, column=1, value="Of the crimes: Calderwick station %s (Abbey Fields), Kingsway station %s (Kingsway)"
                                   % (n0(st["On or near Calderwick (station)"]), n0(st["On or near Kingsway (station)"])))
    ws.cell(row=r, column=1).font = Font(italic=True, color="5B6670")

    r += 2
    ws.cell(row=r, column=1, value="3. Planning figure").font = bold
    r += 1
    for lab, v in (("Calderwick Police arrests (part 1)", M["F"]["F1 the model as filed"]),
                   ("British Transport Police arrests (part 2)", M["B"]["B2 per-type rates with the warrant share"]),
                   ("Detentions 2027-28", M["golden_raw"])):
        ws.cell(row=r, column=1, value=lab)
        ws.cell(row=r, column=5, value=round(v, 1)).number_format = "#,##0.0"
        r += 1
    ws.cell(row=r - 1, column=1).font = ws.cell(row=r - 1, column=5).font = bold
    ws.cell(row=r, column=1, value="Planning figure, nearest 250").font = bold
    c = ws.cell(row=r, column=5, value=M["golden"])
    c.number_format, c.font = "#,##0", bold
    for col, wd in zip("ABCDE", (52, 16, 18, 18, 14)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A4"

    ws2 = wb.create_sheet("Remand holds")
    ws2["A1"] = "Remand holds, defendants charged in 2025-26 and kept in custody for court"
    ws2["A1"].font = Font(bold=True, size=13)
    ws2["A2"] = ("Cell hours from charge to handover to the court escort; held more than 24 hours counts from booking in "
                 "to handover. Co-defendants on a file are counted on their own charge and booking times.")
    header(ws2, 4, ["Suite", "Defendants", "Cell hours, charge to handover", "Held more than 24 hours from booking in"])
    rr = 4
    for s in GD.SUITES:
        rr += 1
        ws2.cell(row=rr, column=1, value=s)
        ws2.cell(row=rr, column=2, value=A[s]["defendants"]).number_format = "#,##0"
        ws2.cell(row=rr, column=3, value=A[s]["cell_hours"]).number_format = "#,##0"
        ws2.cell(row=rr, column=4, value=A[s]["over_24"]).number_format = "#,##0"
    rr += 1
    for j, v in ((1, "Four suites"), (2, sum(A[s]["defendants"] for s in GD.SUITES)),
                 (4, sum(A[s]["over_24"] for s in GD.SUITES))):
        c = ws2.cell(row=rr, column=j, value=v)
        c.font, c.border = bold, top
        if j > 1:
            c.number_format = "#,##0"
    for col, wd in zip("ABCD", (16, 13, 18, 22)):
        ws2.column_dimensions[col].width = wd
    ws2.row_dimensions[4].height = 32

    ws3 = wb.create_sheet("Health assessments")
    ws3["A1"] = "Initial health assessments in 2025-26, by the suite where they took place"
    ws3["A1"].font = Font(bold=True, size=13)
    ws3["A2"] = ("An assessment continued by the incoming clinician at the 07:00 or 19:00 handover is one assessment. "
                 "Eastmoor's team worked at Northgate from 14 August to 2 September 2025, and Wendle Cross's at "
                 "Southfields from 3 November 2025 to 30 January 2026.")
    header(ws3, 4, ["Suite", "Begun 07:00 to 19:00", "Begun 19:00 to 07:00", "2025-26"])
    rr = 4
    for s in GD.SUITES:
        rr += 1
        ws3.cell(row=rr, column=1, value=s)
        ws3.cell(row=rr, column=2, value=B[s]["day"]).number_format = "#,##0"
        ws3.cell(row=rr, column=3, value=B[s]["night"]).number_format = "#,##0"
        ws3.cell(row=rr, column=4, value=B[s]["day"] + B[s]["night"]).number_format = "#,##0"
    rr += 1
    for j, v in ((1, "Four suites"), (2, sum(B[s]["day"] for s in GD.SUITES)), (3, sum(B[s]["night"] for s in GD.SUITES)),
                 (4, sum(B[s]["day"] + B[s]["night"] for s in GD.SUITES))):
        c = ws3.cell(row=rr, column=j, value=v)
        c.font, c.border = bold, top
        if j > 1:
            c.number_format = "#,##0"
    for col, wd in zip("ABCD", (16, 18, 18, 11)):
        ws3.column_dimensions[col].width = wd
    for w in (ws2, ws3):
        w["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        w.merge_cells("A2:D2")
        w.row_dimensions[2].height = 46
        w.freeze_panes = "A5"
    ws4 = wb.create_sheet("Notes")
    ws4["A1"] = "Notes"
    ws4["A1"].font = Font(bold=True, size=12)
    notes = [
        "Prepared by Custody Services Planning on 26 October 2026 for the Tanner Street fit-out contract.",
        "Financial years run from 1 April to 31 March by the date of arrest in UK time (custody demand model, 2.2); the "
        "arrest records give times in UTC.",
        "Part 1 uses the force's arrest records pulled on 6 October 2026. No hub ward reaches the model's 10 per cent cap.",
        "Part 2 uses data.police.uk street-level crime for April 2025 to March 2026, placed in wards with the 2021 LSOA to "
        "ward best-fit lookup. Arrests per recorded crime are the force's 2025-26 arrests in each offence group over its "
        "own street-level crimes of that type; warrant, recall or breach arrests have no crime type and are added at their "
        "share of the force's arrests for crime.",
        "Remand holds use the Criminal Justice Unit's remand files. A file carries the arrest its case started with, which "
        "for a person released and arrested again is the earlier arrest, and co-defendants charged on a file have no file "
        "of their own.",
        "Initial health assessments use Calderwick Healthcare Partners' contact export, whose suite column is the "
        "clinician's team.",
    ]
    for k, t in enumerate(notes, 1):
        ws4.cell(row=k + 2, column=1, value=k).alignment = Alignment(vertical="top")
        c = ws4.cell(row=k + 2, column=2, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws4.column_dimensions["A"].width = 4
    ws4.column_dimensions["B"].width = 100
    for w in wb.worksheets:
        w.page_setup.orientation = "landscape"
        w.page_setup.fitToWidth = 1
        w.page_setup.fitToHeight = 0
        w.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)
    PP.normalize_ooxml(path, PLANNER, PLANNER, "2026-10-21T08:12:00Z", "2026-10-26T15:40:00Z", excel=True)


# ------------------------------------------------------------------------------------------------------------------
def chart_png(R):
    M = R["M"]
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9, "svg.hashsalt": "task137",
                         "axes.edgecolor": "#a7b0b8", "axes.labelcolor": INK, "xtick.color": GREY, "ytick.color": GREY})
    fig, ax = plt.subplots(figsize=(6.4, 3.4), dpi=200)
    years = sorted(M["fy_hub"])
    labels = ["%d-%02d" % (y, (y + 1) % 100) for y in years] + ["2027-28\nplanning figure"]
    own = [M["fy_hub"][y] for y in years]
    x = list(range(len(years))) + [len(years) + 0.6]
    ax.bar(x[:-1], own, width=0.62, color=BLUE, label="Arrests by Calderwick Police officers", zorder=2)
    f_own, f_btp = M["F"]["F1 the model as filed"], M["B"]["B2 per-type rates with the warrant share"]
    ax.bar([x[-1]], [f_own], width=0.62, color=BLUE, alpha=0.6, zorder=2, linewidth=0)
    ax.bar([x[-1]], [f_btp], bottom=[f_own + 12], width=0.62, color=ORANGE, zorder=2,
           label="Arrests by British Transport Police officers at Calderwick and Kingsway stations")
    ax.text(x[-1], f_own + f_btp + 80, n0(M["golden"]), ha="center", va="bottom", color=INK, fontsize=9, fontweight="bold")
    ax.text(x[-1], f_own / 2, n0(f_own), ha="center", va="center", color="white", fontsize=7.5)
    ax.text(x[-1] + 0.36, f_own + f_btp / 2, n0(f_btp), ha="left", va="center", color=INK, fontsize=7.5)
    for xi, v in zip(x[:-1], own):
        ax.text(xi, v + 60, n0(v), ha="center", va="bottom", color=GREY, fontsize=7.5)
    ax.set_xticks(x)
    ax.set_xticklabels(labels)
    ax.set_xlim(-0.6, x[-1] + 0.75)
    ax.set_ylim(0, 4700)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: "{:,.0f}".format(v)))
    ax.set_ylabel("Arrests in the six wards")
    ax.yaxis.grid(True, color="#e3e6e9", linewidth=0.6, zorder=0)
    ax.set_axisbelow(True)
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    ax.axvline(len(years) - 0.2, color="#a7b0b8", linewidth=0.6, linestyle=(0, (2, 2)))
    ax.legend(loc="upper left", frameon=False, fontsize=7.5, bbox_to_anchor=(0.0, 1.2), ncol=1)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def _p(d, text, size=10.5, bold=False, italic=False, color=None, after=6, before=0):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold, r.italic = bold, italic
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    return p


def write_docx(R, path, png):
    M, A, B = R["M"], R["A"], R["B"]
    d = docx.Document()
    sec = d.sections[0]
    sec.page_height, sec.page_width = Mm(297), Mm(210)
    sec.left_margin = sec.right_margin = Mm(22)
    sec.top_margin, sec.bottom_margin = Mm(18), Mm(18)
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10.5)
    _p(d, "Calderwick Police  |  Custody Services Planning", size=9, color="5B6670", after=2)
    _p(d, "Tanner Street custody hub: %s detentions in 2027-28" % n0(M["golden"]), size=15, bold=True, color="1B365D", after=4)
    _p(d, "To: %s, Programme Manager\nFrom: %s, Custody Services Planning\nDate: 26 October 2026" % (PROGRAMME, PLANNER),
       size=9.5, color="5B6670", after=10)
    st = M["btp_hub_by_station"]
    _p(d, "The planning figure for the hub's first year, 1 April 2027 to 31 March 2028, is %s detentions (%s before "
          "rounding to the nearest 250)." % (n0(M["golden"]), "{:,.1f}".format(M["golden_raw"])), bold=True)
    _p(d, "It has two parts. The custody demand model carries the six wards' arrests by our own officers forward from "
          "2025-26 at each ward's capped three-year change, which gives %s for 2027-28 (%s in 2025-26). The second part "
          "is British Transport Police's arrests at Calderwick and Kingsway stations, which sit in Abbey Fields and "
          "Kingsway. BTP has no cells of its own and books the people its officers arrest into the suite serving the "
          "ward, so from April 2027 they come to Tanner Street, but they are in none of our arrest records and the model "
          "never sees them. BTP recorded %s crimes at the two stations in 2025-26 (%s at Calderwick, %s at Kingsway); "
          "at our own arrests per recorded crime of each type, with our share of warrant and recall arrests, that is "
          "%s arrests a year." % (n0(M["F"]["F1 the model as filed"]), n0(M["own_2025"]), n0(M["btp_hub_crimes"]),
                                  n0(st["On or near Calderwick (station)"]), n0(st["On or near Kingsway (station)"]),
                                  n0(M["B"]["B2 per-type rates with the warrant share"])))
    _p(d, "Run on our own records alone the model gives %s to the nearest 250, and the cells and detention officer "
          "posts would be sized for a year without BTP's detainees." % n0(M["default"]))
    d.add_picture(io.BytesIO(png), width=Mm(158))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _p(d, "Arrests in the six wards by financial year from our records, and the 2027-28 planning figure in its two "
          "parts.", size=8.5, italic=True, color="5B6670", after=10)
    _p(d, "The workbook also gives the remand holds and initial health assessments the contract sizes the remand cells "
          "and medical rooms from, suite by suite. Across the four suites, %s defendants charged in 2025-26 were kept in "
          "custody for court, and %s of them were held for more than 24 hours from booking in. There were %s initial "
          "health assessments, %s of them begun at night."
       % (n0(sum(A[s]["defendants"] for s in GD.SUITES)), n0(sum(A[s]["over_24"] for s in GD.SUITES)),
          n0(sum(B[s]["day"] + B[s]["night"] for s in GD.SUITES)), n0(sum(B[s]["night"] for s in GD.SUITES))))
    _p(d, "Financial years run from 1 April to 31 March by the date of arrest in UK time, as the model counts them. The "
          "model's own back-test on the four open suites came within 4 per cent.", size=8.5, color="5B6670", after=8)
    _p(d, PLANNER.split()[0], before=6, after=0)
    ft = sec.footer.paragraphs[0]
    ft.text = "Custody Services Planning  |  Tanner Street planning figure 2027-28"
    ft.runs[0].font.size = Pt(8)
    ft.runs[0].font.color.rgb = RGBColor.from_string("5B6670")
    d.save(path)
    PP.normalize_ooxml(path, PLANNER, PLANNER, "2026-10-24T09:05:00Z", "2026-10-26T15:52:00Z")


def main(tgt, out_dir, R=None):
    R = R or GD.figures(tgt)
    os.makedirs(out_dir, exist_ok=True)
    write_xlsx(R, os.path.join(out_dir, F_XLSX))
    write_docx(R, os.path.join(out_dir, F_DOCX), chart_png(R))
    return R


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else GD.BD.TARGET,
         sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(HERE), "golden"))
