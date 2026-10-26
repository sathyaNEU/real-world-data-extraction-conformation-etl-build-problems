"""The pack's documents for task121: the quality assurance program's measures and definitions and the protocol software
release note (PDF), the centre reference workbook (XLSX), the data dictionary (CSV) and the folder index (TXT)."""
from __future__ import annotations

import csv

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                    # noqa: E402
from reportlab.lib.pagesizes import LETTER                                          # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                     # noqa: E402
from reportlab.lib.units import inch                                                # noqa: E402
from reportlab.pdfbase import pdfmetrics                                            # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                        # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle      # noqa: E402

import common as K   # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("Sans", "LiberationSans-Regular"), ("Sans-B", "LiberationSans-Bold"), ("Sans-I", "LiberationSans-Italic")):
    pdfmetrics.registerFont(TTFont(nm, FONT_DIR + f + ".ttf"))
pdfmetrics.registerFontFamily("Sans", normal="Sans", bold="Sans-B", italic="Sans-I", boldItalic="Sans-B")
INK, GREY, RULE = colors.HexColor("#1b2430"), colors.HexColor("#56616c"), colors.HexColor("#b8c0c8")
NAVY = colors.HexColor("#1f3a5f")
RED = colors.HexColor("#8a2a1f")


def styles(accent=NAVY):
    body = ParagraphStyle("b", fontName="Sans", fontSize=9.8, leading=13.4, textColor=INK, spaceAfter=5)
    return dict(
        body=body,
        h1=ParagraphStyle("h1", fontName="Sans-B", fontSize=14.5, leading=18, textColor=accent, spaceAfter=2),
        h2=ParagraphStyle("h2", fontName="Sans-B", fontSize=10.6, leading=14, textColor=INK, spaceBefore=8,
                          spaceAfter=2, keepWithNext=1),
        meta=ParagraphStyle("m", fontName="Sans", fontSize=8.4, leading=11.2, textColor=GREY, spaceAfter=9),
        cell=ParagraphStyle("c", fontName="Sans", fontSize=8.2, leading=10.4, textColor=INK),
        cellb=ParagraphStyle("cb", fontName="Sans-B", fontSize=8.2, leading=10.4, textColor=INK),
        clause=ParagraphStyle("p", parent=body, leftIndent=28, firstLineIndent=-28),
        bullet=ParagraphStyle("u", parent=body, leftIndent=14, firstLineIndent=-9, spaceAfter=3))


def _frame(left, right):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(GREY)
        canvas.setFont("Sans", 7.4)
        canvas.drawString(0.8 * inch, LETTER[1] - 0.5 * inch, left)
        canvas.drawRightString(LETTER[0] - 0.8 * inch, LETTER[1] - 0.5 * inch, right)
        canvas.drawRightString(LETTER[0] - 0.8 * inch, 0.45 * inch, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _table(S, rows, widths):
    data = [[Paragraph(str(c), S["cellb"] if i == 0 else S["cell"]) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=[w * inch for w in widths], repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, 0), .6, INK), ("LINEBELOW", (0, -1), (-1, -1), .4, RULE),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 1.8),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 1.8)]))
    return t


def _clause(S, num, text):
    return Paragraph("<b>%s</b>&nbsp;&nbsp;%s" % (num, text), S["clause"])


def _doc(path, title, author, subject, left, right):
    return SimpleDocTemplate(path, pagesize=LETTER, leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                             topMargin=0.8 * inch, bottomMargin=0.75 * inch, title=title, author=author,
                             subject=subject, creator=author, producer=author), _frame(left, right)


# ------------------------------------------------------------------------------------------------ QA definitions
def write_definitions(path):
    S = styles()
    title = "Quality Assurance Program: Measures and Definitions"
    doc, fr = _doc(path, title, K.PEOPLE["qa"], "Quality assurance", K.CENTRE, "QA program, effective 1 July 2025")
    st = [Paragraph(title, S["h1"]),
          Paragraph("%s. Effective 1 July 2025. Prepared by %s, Quality Assurance Manager; approved by %s, EMS "
                    "Medical Director, for the %s." % (K.CENTRE, K.PEOPLE["qa"], K.PEOPLE["md"], K.COMMITTEE_NAME),
                    S["meta"]),
          Paragraph("1. The program", S["h2"]),
          _clause(S, "1.1", "The quality assurance program reports the centre's call-taking measures to the EMS "
                            "Quality Committee each month. When a measure moves, the program reports what caused the "
                            "move: the cause is the change without which the measure would have stayed where it was."),
          _clause(S, "1.2", "Months are the centre's calendar months in local time (US Central). A call belongs to "
                            "the month in which it was answered, or for a call never answered, the month in which it "
                            "arrived."),
          Paragraph("2. Triage", S["h2"]),
          _clause(S, "2.1", "A medical call is dispatched hot (lights and sirens) at a determinant of level D or E, "
                            "and cold at a determinant of level O, A, B or C, under the county's EMS response plan."),
          _clause(S, "2.2", "A medical call is under-triaged when it was dispatched cold and the crew assessed a "
                            "patient on it at acuity 1 (critical) or 2 (emergent). The under-triage rate is the "
                            "under-triaged calls over the medical calls on which the crew assessed a patient at "
                            "acuity 1 or 2. A call counts once, at its most acute patient."),
          _clause(S, "2.3", "The crew's acuity is the acuity on the electronic patient care record, matched to the "
                            "call by its incident number."),
          Paragraph("3. Call processing", S["h2"]),
          _clause(S, "3.1", "Call processing time runs from the call being answered to the first ambulance being "
                            "assigned to it."),
          _clause(S, "3.2", "The program reports each month's 90th percentile: the month's medical calls in order of "
                            "their processing time, the time of the call at rank 0.9 times the number of calls, "
                            "rounded up."),
          Paragraph("4. Answering", S["h2"]),
          _clause(S, "4.1", "The centre answers to the national standard: 90 per cent of 911 calls answered within "
                            "15 seconds of arriving at the centre."),
          _clause(S, "4.2", "The share is the month's 911 calls answered within 15 seconds of arriving (15 seconds "
                            "or less) over the month's 911 calls answered. A call abandoned before it was answered is "
                            "reported on its own and is not in the share."),
          Paragraph("5. Telecommunicator CPR", S["h2"]),
          _clause(S, "5.1", "Every cardiac arrest call is reviewed from its recording. Time to first compression runs "
                            "from the call being answered to the first chest compression the call-taker coached."),
          _clause(S, "5.2", "The program reports each month's median time to first compression over the arrests in "
                            "which the call-taker coached compressions, each arrest at its primary review. A second "
                            "review of the same call, for reviewer calibration, is not counted."),
          Paragraph("6. Records", S["h2"]),
          _clause(S, "6.1", "Calls, units and call-takers are as the CAD records them; answering is as the 911 call "
                            "handling equipment records it; acuity is as the patient care record records it; shifts "
                            "are as the scheduling system records the shifts worked.")]
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ------------------------------------------------------------------------------------------------ release note
def write_release(path):
    S = styles(RED)
    title = "Protocol Software Release 13.3"
    doc, fr = _doc(path, title, "QCECD Systems", "Release note", K.CENTRE, "Release 13.3")
    rows = [["Protocol", "Change in 13.3"],
            ["10 Chest pain", "Key questions reordered; new question on pain spreading to the jaw or arm."],
            ["17 Falls", "Determinant suffixes revised for long falls and falls from height; key question on "
                         "anticoagulant medicines."],
            ["26 Sick person", "Key questions on alertness and breathing moved ahead of the chief complaint "
                               "questions."],
            ["28 Stroke", "Stroke diagnostic tool updated to the current scale; time last seen well recorded."],
            ["All protocols", "Breathing-status prompt shown at case entry for every caller."]]
    st = [Paragraph(title, S["h1"]),
          Paragraph("%s, systems team. Issued 20 February 2026." % K.CENTRE, S["meta"]),
          Paragraph("Go-live", S["h2"]),
          Paragraph("Release 13.3 of the protocol software goes live on every call-taking position at 06:00 on "
                    "Tuesday 3 March 2026, replacing release 13.1. Calls open at 06:00 finish on the release they "
                    "started on.", S["body"]),
          Paragraph("Changes for call-takers", S["h2"]),
          _table(S, rows, [1.6, 5.2]),
          Paragraph("Training", S["h2"]),
          Paragraph("Every call-taker completes the two-hour update course before go-live. The course is recorded "
                    "in the training system; supervisors confirm completion at the shift briefing.", S["body"]),
          Paragraph("Support", S["h2"]),
          Paragraph("The vendor's support line is staffed for the first 72 hours after go-live. Report problems to "
                    "the systems team through the usual ticket queue.", S["body"])]
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ------------------------------------------------------------------------------------------------ dictionary
DICTIONARY = [
    ("takers", "employee_no", "The scheduling system's employee number."),
    ("takers", "cad_login", "The login the call-taker uses on the CAD."),
    ("takers", "emd_certified", "Date the call-taker was first certified in emergency medical dispatch."),
    ("takers", "academy_class", "The centre's training academy class, where the call-taker came through it."),
    ("takers", "status", "Employment status at the extract date."),
    ("takers", "status_date", "Date of the status, where it changed in the period."),
    ("roster", "employee_no", "The scheduling system's employee number."),
    ("roster", "shift_date", "Date the shift started."),
    ("roster", "shift", "Day (07:00 to 15:00), Evening (15:00 to 23:00) or Night (23:00 to 07:00)."),
    ("roster", "start", "Scheduled start of the shift worked, local time."),
    ("roster", "end", "Scheduled end of the shift worked, local time."),
    ("roster", "worked_as", "Regular, Swap (worked in a colleague's place) or Overtime."),
    ("calls", "incident_no", "The CAD incident number."),
    ("calls", "call_answered", "Time the call was answered at the call-taker's position, local time, as the CAD "
                               "takes it from the phone system."),
    ("calls", "call_taker", "CAD login of the call-taker who answered and triaged the call."),
    ("calls", "protocol", "Protocol the call was triaged on."),
    ("calls", "chief_complaint", "The protocol's chief complaint."),
    ("calls", "determinant", "Determinant code: protocol, level and suffix."),
    ("calls", "response", "HOT or COLD, from the determinant level under the response plan."),
    ("calls", "protocol_version", "Release of the protocol software the call was triaged on."),
    ("calls", "unit", "The ambulance on the call."),
    ("calls", "unit_assigned", "Time the ambulance on the call was assigned to it, local time."),
    ("calls", "che_call_id", "The phone system's call identifier for the call the incident was opened from."),
    ("epcr", "pcr_number", "Patient care record number."),
    ("epcr", "incident_number", "The CAD incident number as the crew's tablet records it."),
    ("epcr", "unit", "The ambulance whose crew wrote the record."),
    ("epcr", "patient", "Patient number on the incident."),
    ("epcr", "crew_acuity", "Acuity the crew assessed the patient at."),
    ("epcr", "disposition", "What happened to the patient."),
    ("epcr", "pcr_completed", "Time the crew completed the record, local time."),
    ("units", "incident_no", "The CAD incident number."),
    ("units", "unit", "Ambulance."),
    ("units", "assigned", "Time the unit was assigned to the incident, local time."),
    ("units", "enroute", "Time the unit reported en route, local time."),
    ("units", "at_scene", "Time the unit reported at scene, local time."),
    ("units", "cleared", "Time the unit cleared the incident or was taken off it, local time."),
    ("units", "outcome", "Arrived, Diverted (taken off for a higher-priority call) or Cancelled (replaced by a "
                         "closer unit)."),
    ("che", "call_id", "The phone system's call identifier."),
    ("che", "trunk_group", "Trunk group the call arrived on."),
    ("che", "class_of_service", "Class of service delivered with the call."),
    ("che", "ani", "Calling number delivered with the call, masked to its last four digits."),
    ("che", "arrived_utc", "Time the call arrived, UTC."),
    ("che", "answered_utc", "Time a position answered the call, UTC; blank when the caller hung up first."),
    ("che", "released_utc", "Time the call was released, UTC."),
    ("che", "position", "Position that answered the call."),
    ("che", "abandoned", "Y when the caller hung up before the call was answered."),
    ("tcpr", "review_id", "Review number."),
    ("tcpr", "incident_no", "The CAD incident number of the arrest call."),
    ("tcpr", "review_type", "Primary, or Calibration (a second review of the same call by another reviewer)."),
    ("tcpr", "reviewer", "The reviewer's code."),
    ("tcpr", "reviewed_on", "Date of the review."),
    ("tcpr", "recording_started", "Time the logging recorder started the call's recording, local time."),
    ("tcpr", "arrest_recognised_s", "Seconds into the recording at which the call-taker recognised the arrest."),
    ("tcpr", "first_compression_s", "Seconds into the recording at which the first chest compression the "
                                    "call-taker coached was given; blank when none was coached."),
    ("tcpr", "no_tcpr_reason", "Why no compressions were coached, where none were."),
    ("casereview", "case_id", "Case review number."),
    ("casereview", "incident_no", "The CAD incident number reviewed."),
    ("casereview", "call_taker", "CAD login of the call-taker."),
    ("casereview", "reviewed_on", "Date of the review."),
    ("casereview", "case_entry_pct", "Compliance score, case entry questions."),
    ("casereview", "chief_complaint_pct", "Compliance score, chief complaint selection."),
    ("casereview", "key_questions_pct", "Compliance score, key questions."),
    ("casereview", "final_coding_pct", "Compliance score, final coding."),
    ("casereview", "overall_pct", "Overall compliance score, the mean of the four."),
    ("unithours", "service_date", "Date."),
    ("unithours", "unit", "Ambulance."),
    ("unithours", "staffed_hours", "Hours the unit was staffed and in service that day."),
]
SYSTEM = {"takers": "Scheduling system", "roster": "Scheduling system", "calls": "CAD", "epcr": "ePCR",
          "units": "CAD", "che": "911 call handling equipment", "tcpr": "QA review log",
          "casereview": "QA review log", "unithours": "EMS operations"}
NUMERIC = {("roster", "worked_hours"): "decimal", ("tcpr", "arrest_recognised_s"): "integer",
           ("tcpr", "first_compression_s"): "integer", ("epcr", "patient"): "integer",
           ("unithours", "staffed_hours"): "decimal(4,1)"}
DATETIME = {"call_answered", "unit_assigned", "pcr_completed", "assigned", "enroute", "at_scene", "cleared",
            "arrived_utc", "answered_utc", "released_utc", "recording_started", "start", "end"}
DATE = {"emd_certified", "status_date", "shift_date", "reviewed_on", "service_date"}


def _dtype(k, fld):
    if (k, fld) in NUMERIC:
        return NUMERIC[(k, fld)]
    if fld.endswith("_pct"):
        return "integer"
    if fld in DATETIME:
        return "datetime"
    if fld in DATE:
        return "date"
    return "text"


def write_dictionary(path):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(["file_name", "field_name", "source_system", "data_type", "description"])
        for k, fld, desc in DICTIONARY:
            w.writerow([K.FILES[k], fld, SYSTEM[k], _dtype(k, fld), desc])
        for sheet, desc in (("Trunk groups", "One row per trunk group the centre's calls arrive on."),
                            ("Units", "One row per ambulance: type and station."),
                            ("Protocols", "One row per protocol the centre triages medical calls on."),
                            ("Positions", "One row per position that answers calls.")):
            w.writerow([K.FILES["reference"], "sheet " + sheet, "Reference workbook", "sheet", desc])


# ------------------------------------------------------------------------------------------------ reference workbook
def _style_sheet(sh, widths=None, fill="1F3A5F"):
    from openpyxl.styles import Alignment, Font, PatternFill
    for cell in sh[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sh.freeze_panes = "A2"
    for col in sh.columns:
        letter_ = col[0].column_letter
        if widths and letter_ in widths:
            sh.column_dimensions[letter_].width = widths[letter_]
        else:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:200])
            sh.column_dimensions[letter_].width = min(60, max(8, width + 2))


def write_reference(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Trunk groups"
    ws.append(["trunk_group", "description", "circuits"])
    for tg, desc, _ in K.TRUNKS:
        ws.append([tg, desc, {"TG01": 6, "TG02": 12, "TG03": 4, "TG05": 2, "TG07": 4}[tg]])
    _style_sheet(ws, {"B": 70})
    wu = wb.create_sheet("Units")
    wu.append(["unit", "type", "station"])
    for i, u in enumerate(K.UNITS):
        wu.append([u, "ALS" if i % 3 else "BLS", "Station %d" % (1 + i // 2)])
    _style_sheet(wu)
    wp = wb.create_sheet("Protocols")
    wp.append(["protocol", "chief_complaint"])
    for p in K.PROTOCOLS:
        wp.append([p[0], p[1]])
    _style_sheet(wp)
    wq = wb.create_sheet("Positions")
    wq.append(["position", "room"])
    for i in range(1, 21):
        wq.append(["P%02d" % i, "Call floor" if i <= 12 else ("Training room (overflow)" if i <= 14 else
                                                             "Dispatch floor")])
    _style_sheet(wq)
    wb.save(path)


def write_index(path, items):
    out = ["%s: medical call under-triage review, files for the %s" % (K.CENTRE, K.COMMITTEE_NAME),
           "Assembled %s by %s, Quality Assurance." % (K.ASSEMBLED.strftime("%B %d, %Y").replace(" 0", " "),
                                                       K.PEOPLE["qa"]), ""]
    for name, desc in items:
        out.append(name)
        out.append("    " + desc)
    with open(path, "w", encoding="utf-8", newline="\r\n") as f:
        f.write("\n".join(out) + "\n")
