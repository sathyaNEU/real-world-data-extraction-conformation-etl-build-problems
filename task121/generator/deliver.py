"""The goldens from the shipped pack: the cause workbook (the cause, the rates before and from March, and the month by
month measures the committee asked for) and the chart of each month's under-triage rate with the date the cause began;
and generator/figures.json, every figure the write-up states."""
import datetime as dt
import json
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

import analysis as A   # noqa: E402
import asks as AK      # noqa: E402
import common as K     # noqa: E402

INK, INK2, GRID, SURF = "#14202e", "#52606d", "#e2e6ea", "#fbfbfa"
BLUE, ORANGE, GREY = "#2f6db5", "#d9622b", "#9aa5b1"
CAUSE = ("the quick returns of the 2026 rotation: from 2 March 2026 about one shift in six began less than eleven "
         "hours after the call-taker's previous shift ended, and the calls taken on those shifts were sent cold at "
         "three times the rate of the rest")


def ym_key(ym):
    return "%d-%02d" % ym


def compute(target):
    G = A.golden(target)
    rg = A.rungs(target, G)
    ak = AK.golden(target)
    fig = {
        "call": {"cause": "quick returns (the 2026 rotation)", "began": "2026-03-02",
                 "statement": "The rise came from the 2026 rotation's quick returns: shifts begun less than eleven "
                              "hours after the call-taker's previous shift ended, not from protocol release 13.3."},
        "rates": {"before": K.r1(G["before"]), "from_march": K.r1(G["after"])},
        "components": {
            "quick_return_share_before": K.r1(G["qr_share_before"]),
            "quick_return_share_from_march": K.r1(G["qr_share_after"]),
            "from_march_on_quick_return_shifts": K.r1(G["after_qr"]),
            "from_march_on_rested_shifts": K.r1(G["after_rested"]),
            "rested_shifts_before": K.r1(G["before_rested"]),
        },
        "steps": {
            "release_changed_protocols": [K.r1(x) for x in rg["Rp_changed"]],
            "release_unchanged_protocols": [K.r1(x) for x in rg["Rp_unchanged"]],
            "veterans": [K.r1(x) for x in rg["Rg_veterans"]],
            "standardised_from_march": K.r1(rg["Rm_standardised"][1]),
        },
        "monthly_under_triage": {ym_key(ym): K.r1(G["monthly"][ym]) for ym in K.MONTHS},
        "asks": {
            "A": {"p90_processing_s": {ym_key(ym): ak["A"][ym] for ym in K.MONTHS}},
            "B": {"median_first_compression_s": {ym_key(ym): ak["B"][ym] for ym in K.MONTHS}},
            "C": {"answered_within_15s_pct": {ym_key(ym): ak["C"][ym] for ym in K.MONTHS}},
        },
        "threshold": {"quick_return_hours": K.QUICK_RETURN_HOURS, "answer_seconds": K.STANDARD_ANSWER},
        "parameters": {"percentile_rank": 0.9, "qa_clause_under_triage": 2.2, "emergency_number": 911,
                       "release": 13.3},
    }
    return fig, dict(G=G, rungs=rg, asks=ak)


def write_workbook(path, fig):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cause"
    hdr = PatternFill("solid", fgColor="1F3A5F")
    ws.append(["Under-triaged medical calls: what caused the rise from March 2026"])
    ws["A1"].font = Font(bold=True, size=13, color="1F3A5F")
    ws.append([])
    ws.append(["Cause", fig["call"]["statement"]])
    ws.append(["Began", "2 March 2026, the first block of the 2026 rotation"])
    ws.append([])
    ws.append(["Under-triage rate", "Per cent of critical or emergent calls dispatched cold"])
    ws.append(["September 2025 to February 2026", fig["rates"]["before"]])
    ws.append(["March to August 2026", fig["rates"]["from_march"]])
    ws.append([])
    cp = fig["components"]
    ws.append(["Shifts begun under 11 hours after the previous shift", "Per cent of shifts worked"])
    ws.append(["September 2025 to February 2026", cp["quick_return_share_before"]])
    ws.append(["March to August 2026", cp["quick_return_share_from_march"]])
    ws.append([])
    ws.append(["Under-triage rate by the call-taker's rest before the shift", "Per cent"])
    ws.append(["Rested shifts, September 2025 to February 2026", cp["rested_shifts_before"]])
    ws.append(["Rested shifts, March to August 2026", cp["from_march_on_rested_shifts"]])
    ws.append(["Quick-return shifts, March to August 2026", cp["from_march_on_quick_return_shifts"]])
    for r in (3, 6, 10, 14):
        ws.cell(r, 1).font = Font(bold=True)
    for r in (7, 8, 11, 12, 15, 16, 17):
        ws.cell(r, 2).number_format = "0.0"
        ws.cell(r, 2).alignment = Alignment(horizontal="right")
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 96
    ws["B3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[3].height = 32
    wm = wb.create_sheet("Monthly")
    wm.append(["Month", "Under-triage rate (%)", "Call processing, 90th percentile (s)",
               "Time to first compression, median (s)", "911 calls answered within 15 s (%)"])
    for ym in K.MONTHS:
        k = ym_key(ym)
        wm.append([dt.date(ym[0], ym[1], 1), fig["monthly_under_triage"][k], fig["asks"]["A"]["p90_processing_s"][k],
                   fig["asks"]["B"]["median_first_compression_s"][k], fig["asks"]["C"]["answered_within_15s_pct"][k]])
    for cell in wm[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    wm.row_dimensions[1].height = 44
    for row in wm.iter_rows(min_row=2):
        row[0].number_format = "mmm yyyy"
        row[1].number_format = "0.0"
        row[2].number_format = "0"
        row[3].number_format = "0"
        row[4].number_format = "0.0"
    for col, w in zip("ABCDE", (11, 14, 18, 18, 16)):
        wm.column_dimensions[col].width = w
    wm.freeze_panes = "B2"
    n = wb.create_sheet("Notes")
    for line in ["Under-triage: a medical call dispatched cold on which the crew assessed a patient at acuity 1 or 2, "
                 "over the calls on which the crew assessed a patient at acuity 1 or 2, each call at its most acute "
                 "patient (QA program 2.2).",
                 "Rest before a shift: from the end of the call-taker's previous shift worked to the start of the shift, "
                 "from the scheduling system's shifts worked; a quick return is under 11 hours.",
                 "Call processing: answer to the first ambulance assigned, from the CAD's unit history, each call "
                 "once however many incidents carry it (QA program 3).",
                 "Time to first compression: answer to the first coached compression, the recording's seconds taken "
                 "from the start of the recording, each arrest at its primary review; reviewer QA7 keys times as "
                 "minutes and seconds without the colon (QA program 5).",
                 "Answering: 911 trunk groups only (TG01, TG02, TG03 and TG05), calls answered, a call passed "
                 "between positions counted once, by the month answered in local time (QA program 4).",
                 "Sources: CAD, ePCR, scheduling system, 911 call handling equipment and QA review logs, 1 September "
                 "2025 to 31 August 2026."]:
        n.append([line])
    n.column_dimensions["A"].width = 130
    wb.save(path)


def write_chart(path, fig, res):
    G = res["G"]
    xs = [dt.date(y, m, 15) for y, m in K.MONTHS]
    total = [G["monthly"][ym] for ym in K.MONTHS]
    rested = [G["monthly_rested"][ym] for ym in K.MONTHS]
    f, ax = plt.subplots(figsize=(9.0, 5.2), dpi=150)
    f.patch.set_facecolor(SURF)
    ax.set_facecolor(SURF)
    ax.plot(xs, total, color=BLUE, linewidth=2.2, marker="o", markersize=5, zorder=3, label="All critical or emergent calls")
    ax.plot(xs, rested, color=GREY, linewidth=1.6, linestyle=(0, (4, 2)), marker="o", markersize=3.5, zorder=2,
            label="Calls on shifts begun after 11 hours' rest or more")
    cut = dt.date(2026, 3, 2)
    ax.axvline(cut, color=ORANGE, linewidth=1.6, zorder=1)
    ax.text(cut + dt.timedelta(days=3), 21.2, "2 March 2026: the 2026 rotation\nbegins; quick returns start",
            color=ORANGE, fontsize=8.6, va="top")
    ax.set_ylim(0, 22)
    ax.set_ylabel("Under-triage rate (%)", fontsize=9, color=INK2)
    ax.yaxis.grid(True, color=GRID, linewidth=0.8, zorder=0)
    ax.tick_params(axis="both", length=0, colors=INK2, labelsize=8.6)
    import matplotlib.dates as mdates
    ax.xaxis.set_major_locator(mdates.MonthLocator())
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%b\n%Y"))
    for sp in ax.spines.values():
        sp.set_visible(False)
    ax.legend(loc="lower left", frameon=False, fontsize=8.4)
    f.suptitle("Under-triage rose with the 2026 rotation's quick returns, not across the board", x=0.02, ha="left",
               fontsize=12, color=INK, fontweight="bold")
    f.text(0.02, 0.905, "Medical calls dispatched cold among those the crew found critical or emergent, by month, "
           "September 2025 to August 2026", fontsize=8.6, color=INK2)
    f.text(0.02, 0.012, "Source: %s CAD, ePCR and scheduling system. Protocol release 13.3 went live on 3 March "
           "2026." % K.SHORT, fontsize=7.2, color=INK2)
    f.tight_layout(rect=(0, 0.03, 1, 0.89))
    f.savefig(path, dpi=150, facecolor=SURF, metadata={"Software": None})
    plt.close(f)


def write(target, golden):
    os.makedirs(golden, exist_ok=True)
    fig, res = compute(target)
    write_workbook(os.path.join(golden, K.GOLDEN), fig)
    write_chart(os.path.join(golden, K.CHART), fig, res)
    return fig, res


def write_figures(path, fig):
    with open(path, "w") as f:
        json.dump(fig, f, indent=1, sort_keys=True)
        f.write("\n")
