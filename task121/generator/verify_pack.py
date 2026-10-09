#!/usr/bin/env python3
"""Independent verifier for task121. Reads only the shipped files under <task>/target and recomputes,
on its own code path, every rung, every rival-killer, every calibration outcome and every graded
figure. Imports nothing from the generator.

    python3 -I verify_pack.py <task_dir> [--json out.json]
"""
import argparse
import json
import os
import re
import sys
import zipfile

import numpy as np
import pandas as pd

WEEKS = {"B0": "2026-07-27", "B1": "2026-08-03", "B2": "2026-08-10", "B3": "2026-08-17", "B4": "2026-08-24",
         "W1": "2026-08-31", "W2": "2026-09-07", "W3": "2026-09-14", "W4": "2026-09-21"}
BASE = ["B1", "B2", "B3", "B4"]
REV = ["W1", "W2", "W3", "W4"]
STEPS = ["basket", "contact", "address", "delivery", "payment", "confirmation"]
STEPFIX = {"basket": "F2", "contact": "F4", "address": "F1", "delivery": "F5", "payment": "F3"}
POPFIX = {"addr": "F1", "newfan": "F2", "card": "F3", "acct": "F4", "mixed": "F5"}


def find(tgt, prefix, ext):
    m = [f for f in os.listdir(tgt) if f.startswith(prefix) and f.endswith(ext)]
    assert len(m) == 1, (prefix, m)
    return os.path.join(tgt, m[0])


def week_label(ts):
    edges = pd.to_datetime(list(WEEKS.values()))
    names = np.array(list(WEEKS.keys()))
    pos = np.searchsorted(edges.values, ts.values, side="right") - 1
    return names[pos]


def load(tgt):
    d = {}
    d["s"] = pd.read_parquet(find(tgt, "basket_sessions", ".parquet"))
    d["edge"] = pd.read_csv(find(tgt, "edge_bot_verdicts", ".csv"))
    d["tok"] = pd.concat([pd.read_csv(find(tgt, "cdm_token_billing_2026-08", ".csv")),
                          pd.read_csv(find(tgt, "cdm_token_billing_2026-09", ".csv"))])
    d["prof"] = pd.read_csv(find(tgt, "loyalty_profiles", ".csv"), dtype={"club_member_no": "Int64"})
    d["promo"] = pd.read_csv(find(tgt, "promo_redemptions", ".csv"), dtype={"account_id": str})
    d["addr"] = pd.read_csv(find(tgt, "address_book_history", ".csv"), dtype={"postcode": str})
    d["flags"] = json.load(open(find(tgt, "checkout_flags_export", ".json"), encoding="utf-8"))
    d["cards"] = pd.read_csv(find(tgt, "saved_cards_", ".csv"))
    d["cev"] = pd.read_csv(find(tgt, "saved_card_events", ".csv"))
    d["auth"] = pd.read_csv(find(tgt, "tagus_3ds_log", ".csv"))
    d["fin"] = pd.read_csv(find(tgt, "finance_order_export", ".csv"))
    d["cat"] = pd.read_csv(find(tgt, "catalogue_status_history", ".csv"))
    return d


def classify(d):
    s = d["s"].copy()
    s["ts"] = pd.to_datetime(s.started_at)
    s = s[~s.session_id.isin(set(d["edge"].session_id))].copy()
    s["wk"] = week_label(s.ts)
    s["ordered"] = s.order_id.notna()
    s["is_club"] = s.landing_url.str.contains("mpt=", regex=False)
    tok = s.landing_url.str.extract(r"[?&]mpt=([^&]+)")[0]
    mem = tok.map(dict(zip(d["tok"].token, d["tok"].member_no)))
    pr = d["prof"].dropna(subset=["club_member_no"])
    owner = dict(zip(pr.club_member_no.astype(int), pr.account_id))
    # last season's members' code: SOCIO and the seven-digit member number, redeemed on an account
    pm = d["promo"].dropna(subset=["account_id"])
    code = pm.promo_code.str.extract(r"^SOCIO(\d{7})$")[0]
    for m, a in zip(code, pm.account_id):
        if isinstance(m, str):
            owner.setdefault(int(m), a)
    s["owner"] = [owner.get(int(m)) if pd.notna(m) else None for m in mem]
    s["acct"] = s.is_club & s.owner.notna()
    s["newfan"] = s.is_club & s.owner.isna()
    s["vt"] = np.select([s.signed_in, s.is_club, s.new_visitor], ["SI", "CLUB", "NV"], "RG")
    # flag cohort at the session, from the current assignment and the logged moves
    fl = d["flags"]
    cur = pd.Series({x["account_id"]: x["cohort"] for x in fl["assignments"]})
    moves = pd.DataFrame(fl["assignment_moves"])
    moves["at"] = pd.to_datetime(moves.moved_at.str[:19])
    enabled = {x["cohort"]: pd.Timestamp(x["enabled_at"][:19]) for x in fl["rollout"]}
    # latest move at or before the session; before an account's first move, that move's origin
    moves = moves.sort_values(["at", "account_id"])
    first_from = moves.groupby("account_id").from_cohort.first()
    si = s[s.signed_in][["account_id", "ts"]].reset_index().sort_values("ts")
    m = pd.merge_asof(si, moves[["account_id", "at", "to_cohort"]].rename(columns={"at": "ts_m"}),
                      left_on="ts", right_on="ts_m", by="account_id", direction="backward")
    c = m.to_cohort.where(m.to_cohort.notna(), m.account_id.map(first_from))
    c = c.where(c.notna(), m.account_id.map(cur))
    s["coh"] = 0
    s.loc[m["index"].to_numpy(), "coh"] = c.astype(int).to_numpy()
    s["live"] = [k > 0 and t >= enabled[k] for k, t in zip(s.coh, s.ts)]
    # default address at the session; address book times are UTC, sessions Lisbon (UTC+1 in summer)
    ab = d["addr"][d["addr"].is_default].copy()
    ab["lt"] = (pd.to_datetime(ab.changed_at) + pd.Timedelta(hours=1)).astype("datetime64[ns]")
    ab = ab.sort_values("lt")
    sig = s[s.signed_in][["session_id", "account_id", "ts"]].copy()
    sig["ts"] = sig.ts.astype("datetime64[ns]")
    sig = sig.sort_values("ts")
    j = pd.merge_asof(sig, ab[["account_id", "lt", "postcode"]], left_on="ts", right_on="lt", by="account_id")
    pc = s.session_id.map(dict(zip(j.session_id, j.postcode))).fillna("")
    s["addr"] = s.signed_in & s.live & pc.str.match(r"^\d{4}$")
    # default card at the session
    cards = d["cards"]
    issuer = dict(zip(cards.card_id, cards.issuer))
    dflt = dict(zip(cards.account_id[cards.is_default], cards.card_id[cards.is_default]))
    sd = d["cev"][d["cev"].event == "set_default"]
    sw = {a: (pd.Timestamp(t), p) for a, t, p in zip(sd.account_id, sd["at"], sd.previous_default_card_id)}
    iss = []
    for a, t, si in zip(s.account_id, s.ts, s.signed_in):
        if not si:
            iss.append(None)
            continue
        cid = sw[a][1] if (a in sw and t < sw[a][0]) else dflt[a]
        iss.append(issuer[cid])
    s["card"] = s.signed_in & pd.Series(iss, index=s.index).isin(["Bankora", "Finvo"])
    # pre-order status at the session
    cat = d["cat"].copy()
    cat["vf"] = pd.to_datetime(cat.valid_from)
    st = {k: g.sort_values("vf")[["vf", "status"]].values.tolist() for k, g in cat.groupby("sku")}

    def stat(sku, t):
        out = None
        for vf, x in st[sku]:
            if vf <= t:
                out = x
        return out
    s["mixed"] = [len({stat(k, t) for k in b.split(";")} & {"pre_order", "in_stock"}) == 2
                  for b, t in zip(s.basket_skus, s.ts)]
    return s


def own_rates(s):
    b = s[s.wk.isin(BASE)]
    si = b[b.signed_in]
    r = {}
    for p, col in (("addr", "account_id"), ("card", "account_id"), ("acct", "owner")):
        who = set(s.loc[s.wk.isin(REV) & s[p], col].dropna())
        r[p] = si.loc[si.account_id.isin(who), "ordered"].mean()
    r["newfan"] = b.loc[b.vt == "NV", "ordered"].mean()
    r["mixed"] = b.loc[b.mixed, "ordered"].mean()
    r["store"] = b.ordered.mean()
    for v in ("SI", "NV", "RG"):
        r[v] = b.loc[b.vt == v, "ordered"].mean()
    return r


def weekly_loss(s, rate):
    out = {}
    for p in POPFIX:
        g = s[s[p]].groupby("wk").ordered.agg(["size", "sum"])
        out[p] = {w: float(g.loc[w, "size"] * rate[p] - g.loc[w, "sum"]) if w in g.index else 0.0
                  for w in list(WEEKS)[1:]}
    return out


def order_values(d, s):
    f = d["fin"].copy()
    f["ex"] = pd.to_datetime(f.exported_at)
    f = f[f.ex == f.groupby("order_id").ex.transform("max")]
    f = f.drop_duplicates("order_id")
    net = np.where(f.ex < pd.Timestamp("2026-08-10"), f.order_total_eur / 1.23, f.order_total_eur)
    val = dict(zip(f.order_id, net))
    b = s[s.wk.isin(BASE) & s.ordered]
    out = {}
    for p, col in (("addr", "account_id"), ("card", "account_id"), ("acct", "owner")):
        who = set(s.loc[s.wk.isin(REV) & s[p], col].dropna())
        out[p] = b.loc[b.signed_in & b.account_id.isin(who), "order_id"].map(val).mean()
    out["newfan"] = b.loc[b.vt == "NV", "order_id"].map(val).mean()
    out["mixed"] = b.loc[b.mixed, "order_id"].map(val).mean()
    return out


def challenge(d, s):
    a = d["auth"].drop_duplicates("attempt_ref")
    out = {}
    for p in POPFIX:
        ids = set(s.session_id[(s.wk == "W4") & s[p]])
        x = a[a.session_id.isin(ids)]
        out[p] = 100.0 * x.three_ds_status.isin(["C", "D"]).sum() / len(x)
    return out


def stage(s, exclude_club=False):
    x = s[~s.is_club] if exclude_club else s
    i = x.furthest_step.map({k: n for n, k in enumerate(STEPS)})
    b, r = x.wk.isin(BASE), x.wk.isin(REV)
    out = {}
    for n, k in enumerate(STEPS[:5]):
        rate = (i[b] == n).sum() / (i[b] >= n).sum()
        out[STEPFIX[k]] = float((i[r] == n).sum() - (i[r] >= n).sum() * rate)
    return out


def top(d):
    o = sorted(d, key=lambda k: -d[k])
    return o[0], o[1], d[o[0]] / d[o[1]] if d[o[1]] > 0 else float("inf"), o


def read_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def closeout(path):
    sh = read_xlsx(path)
    rows = sh["Weekly"]
    df = pd.DataFrame(rows[1:], columns=rows[0])
    booked = {r[0]: r[1] for r in sh["Close-out"] if r and r[0]}
    pre = df[(~df.campaign_week.astype(bool)) & (df.channel == "rest")]
    rt = pre.groupby("visitor_type").orders.sum() / pre.groupby("visitor_type").basket_sessions.sum()
    rs = pre.orders.sum() / pre.basket_sessions.sum()
    p = df[df.channel == "partner"]
    vt = float((p.basket_sessions * p.visitor_type.map(rt)).sum() - p.orders.sum())
    return {"booked": booked["Orders lost to the campaign"], "visitor_type": vt,
            "store_rate": float(p.basket_sessions.sum() * rs - p.orders.sum()),
            "all_new": float(p.basket_sessions.sum() * rt["new_visitor"] - p.orders.sum())}


def releases(path):
    sh = read_xlsx(path)
    filed = {r[0]: r[6] for r in sh["Releases"][1:]}
    out = {}
    for name, rows in sh.items():
        if not name.endswith(" cohorts"):
            continue
        k = name.split()[0]
        df = pd.DataFrame(rows[1:], columns=rows[0])
        diffs = []
        for w, g in df.groupby("week_start"):
            on, off = g[g.flag_on.astype(bool)], g[~g.flag_on.astype(bool)]
            if len(on) and len(off):
                diffs.append(on.orders.sum() / on.basket_sessions.sum() - off.orders.sum() / off.basket_sessions.sum())
        ws = sorted(df.week_start.unique())
        pre, post = df[df.week_start.isin(ws[:4])], df[df.week_start.isin(ws[4:])]
        cal = post.orders.sum() / post.basket_sessions.sum() - pre.orders.sum() / pre.basket_sessions.sum()
        out[k] = {"filed": filed[k], "cohort": 100 * float(np.mean(diffs)), "calendar": 100 * float(cal)}
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("task")
    ap.add_argument("--json")
    a = ap.parse_args()
    tgt = os.path.join(a.task, "target")
    d = load(tgt)
    s = classify(d)
    rate = own_rates(s)
    L = weekly_loss(s, rate)
    w4 = {p: L[p]["W4"] for p in POPFIX}
    by_fix = {POPFIX[p]: v for p, v in w4.items()}
    call, runner, margin, _ = top(by_fix)
    val = order_values(d, s)
    ch = challenge(d, s)
    out = {"call": call, "runner_up": runner, "call_w4": by_fix[call], "runner_up_w4": by_fix[runner],
           "gap": by_fix[call] - by_fix[runner],
           "L1": {POPFIX[p]: sum(L[p][w] for w in REV) for p in POPFIX},
           "L2": {POPFIX[p]: L[p]["W1"] for p in POPFIX},
           "L3": {POPFIX[p]: L[p]["W4"] * val[p] for p in POPFIX},
           "L4": {POPFIX[p]: ch[p] for p in POPFIX},
           "weekly": {POPFIX[p]: L[p] for p in POPFIX}}
    # cohort sheet: signed-in sessions over the four review weeks by the cohort at the session
    cs = s[s.wk.isin(REV) & s.signed_in].groupby("coh").ordered.agg(["size", "mean"])
    out["cohorts"] = {str(int(k)): [int(r["size"]), 100 * float(r["mean"])] for k, r in cs.iterrows()}
    # the ladder
    W4 = s[s.wk == "W4"]
    sh = lambda m, r: float(m.sum() * r - W4.ordered[m].sum())
    R = {"R0": stage(s),
         "R1": {"F1": w4["addr"], "F2": sh(W4.is_club, rate["store"]), "F3": sh(W4.card, rate["store"]),
                "F4": sh(W4.vt == "RG", rate["store"]), "F5": sh(W4.mixed, rate["store"])},
         "R2": {"F1": w4["addr"], "F2": sh(W4.is_club, rate["NV"]), "F3": sh(W4.card, rate["SI"]),
                "F4": sh(W4.vt == "RG", rate["RG"]),
                "F5": sum(sh(W4.mixed & (W4.vt == v), rate[v]) for v in ("SI", "NV", "RG"))},
         "R3": {"F1": w4["addr"], "F2": sh(W4.is_club, rate["NV"]), "F3": w4["card"],
                "F4": sh(W4.vt == "RG", rate["RG"]), "F5": w4["mixed"]},
         "R4": by_fix}
    out["rungs"] = {}
    for k, v in R.items():
        l1, l2, m, o = top(v)
        out["rungs"][k] = {"leader": l1, "margin": m, "F4_rank": o.index("F4") + 1, "figures": v}
    # rival-killers
    out["killers"] = {
        "R0_check_drained_w4": w4["addr"],
        "R1_closeouts": {k: closeout(find(tgt, k, ".xlsx")) for k in ("closeout_festival", "closeout_kaiju")},
        "R2_mixed_own_w4": w4["mixed"],
        "R3_chain_share_w4": float(W4.acct[W4.is_club].mean()),
        "stage_without_club": top(stage(s, True))[0],
        "four_week_leader": top(out["L1"])[0],
        "p2_at_store_rate": sh(W4.newfan, rate["store"]),
        "p4_at_store_rate": sh(W4.acct, rate["store"]),
    }
    out["releases"] = releases(find(tgt, "release_log", ".xlsx"))
    # twins: cohorts 6 and 8 over their first two live weeks
    tw = {}
    for k in (6, 8):
        x = s[s.signed_in & (s.coh == k) & s.wk.isin(["W1", "W2"])]
        tw[k] = float(len(x) * rate["SI"] - x.ordered.sum())
    out["twins"] = tw
    # referee: the provider's totals
    import pypdf
    txt = "".join(p.extract_text() for p in pypdf.PdfReader(find(tgt, "tagus_3ds_report", ".pdf")).pages)
    m = re.search(r"Total\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)", txt)
    tot = [int(x.replace(",", "")) for x in m.groups()]
    au = d["auth"].drop_duplicates("attempt_ref")
    out["referee_ties"] = tot[0] == len(au) and tot[2] == int(au.three_ds_status.isin(["C", "D"]).sum())
    # context artifact
    sh_ = read_xlsx(find(tgt, "weekly_trading_dashboard", ".xlsx"))["Weekly"]
    dash = [r[:4] for r in sh_[5:] if r and r[0]]
    calc = []
    allw = d["s"].copy()
    allw["wk"] = week_label(pd.to_datetime(allw.started_at))
    allw["ordered"] = allw.order_id.notna()
    for wname, start in WEEKS.items():
        x = allw[allw.wk == wname]  # as published each Monday, before that week's edge verdicts
        calc.append([start, len(x), int(x.ordered.sum()), round(100 * x.ordered.mean(), 2)])
    out["dashboard_reproduces"] = [list(r) for r in dash] == calc
    out["sessions"] = int(len(d["s"]))
    # structural verdicts
    want = {"R0": "F1", "R1": "F2", "R2": "F5", "R3": "F3", "R4": "F4"}
    checks = {f"{k} names {v}": out["rungs"][k]["leader"] == v and out["rungs"][k]["margin"] >= 1.2
              for k, v in want.items()}
    checks["F4 4th or 5th on the natural pipeline"] = out["rungs"]["R0"]["F4_rank"] in (4, 5)
    checks["close-outs reproduce on visitor type only"] = all(
        abs(v["visitor_type"] - v["booked"]) < 0.5 and v["store_rate"] > 60 and v["all_new"] < -30
        for v in out["killers"]["R1_closeouts"].values())
    checks["releases: cohort reproduces, calendar misses"] = all(
        abs(v["filed"] - v["cohort"]) < 0.01 and abs(v["calendar"] - v["cohort"]) >= 0.3 * abs(v["cohort"])
        for v in out["releases"].values())
    checks["twins two to one"] = tw[6] >= 2 * tw[8]
    checks["referee ties"] = out["referee_ties"]
    checks["dashboard reproduces"] = out["dashboard_reproduces"]
    checks["call F4 over F3"] = call == "F4" and runner == "F3" and margin >= 1.2
    out["checks"] = checks
    bad = [k for k, v in checks.items() if not v]
    print(f"verify_pack: {len(checks) - len(bad)} of {len(checks)} structural checks pass; call {call} "
          f"({by_fix[call]:.2f}) over {runner} ({by_fix[runner]:.2f}), gap {by_fix[call] - by_fix[runner]:.2f}")
    if a.json:
        with open(a.json, "w") as f:
            json.dump(out, f, indent=1, default=float)
    if bad:
        print("FAILED:", bad)
        sys.exit(1)


if __name__ == "__main__":
    main()
