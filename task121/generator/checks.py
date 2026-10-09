"""Every assertion the build carries. Each check fails the build at the moment a parameter or a
record change breaks it. Counted; the count is printed and recorded."""
import datetime as dt
import itertools
import json
import os
import re
import subprocess
import sys
import zipfile

import numpy as np
import pandas as pd

import analysis as N
import calib as C
import params as P

UNIT = {"orders": 1.0, "eur": 100.0, "share": 0.1, "conv": 0.01}
MID = {"orders": 0.2, "eur": 15.0, "share": 0.015, "conv": 0.002}


class Chk:
    def __init__(self):
        self.n = 0
        self.log = []

    def ok(self, cond, msg):
        if not cond:
            raise AssertionError(msg)
        self.n += 1
        self.log.append(msg)


def dist(x, kind):
    u = UNIT[kind]
    f = x / u - np.floor(x / u)
    return abs(f - 0.5) * u


def rnd(x, kind):
    u = UNIT[kind]
    return int(np.floor(x / u + 0.5)) if kind != "share" and kind != "conv" else round(np.floor(x / u + 0.5) * u, 2)


def same(a, b, kind):
    return rnd(a, kind) == rnd(b, kind)


def ask_figs(g, coh):
    """Every graded ask figure with its unit: A1 L1/L2, A2 L3, A3 L4, A4 cohorts."""
    out = {}
    for p in N.POPS:
        out[("A1", "L1", p)] = (g["L1"][p], "orders")
        out[("A1", "L2", p)] = (g["L2"][p], "orders")
        out[("A2", "L3", p)] = (g["L3"][p], "eur")
        out[("A3", "L4", p)] = (g["L4"][p], "share")
    for c, (n, v) in coh.items():
        out[("A4", "n", c)] = (float(n), "orders")
        out[("A4", "conv", c)] = (v, "conv")
    return out


def moved(fa, fb, ask=None):
    """Keys whose figure falls in a different rounding bin."""
    return sorted(k for k in fa if (ask is None or k[0] == ask) and k in fb and not same(fa[k][0], fb[k][0], fa[k][1]))


# --------------------------------------------------------------------------- main ladder

def ladder(c, ctx, g, s):
    pk = ctx["pk"]
    R = N.rungs(s, g["base"], g["L"])
    want = {"R0": "F1", "R1": "F2", "R2": "F5", "R3": "F3", "R4": "F4"}
    rec = {}
    for r, d in R.items():
        o = N.rank(d)
        c.ok(o[0] == want[r], f"{r} names {want[r]} (got {o[0]})")
        m = d[o[0]] / d[o[1]] if d[o[1]] > 0 else float("inf")
        c.ok(m >= 1.2, f"{r} margin {m:.2f}x >= 1.2")
        rec[r] = {"figures": {k: round(v, 2) for k, v in d.items()}, "leader": o[0], "margin": round(m, 2),
                  "F4_rank": o.index("F4") + 1}
    c.ok(rec["R0"]["F4_rank"] in (4, 5), f"F4 ranks {rec['R0']['F4_rank']} of 5 on the natural pipeline")
    c.ok(all(rec[r]["F4_rank"] != 1 for r in ("R0", "R1", "R2", "R3")), "F4 leads no intermediate rung")
    c.ok(sum(rec[r]["F4_rank"] == 2 for r in ("R0", "R1", "R2", "R3")) <= 1, "F4 second on at most one rung")
    names = [want[r] for r in ("R0", "R1", "R2", "R3", "R4")]
    c.ok(all(names[i] != names[i + 1] for i in range(4)), "no adjacent rungs share a leader")
    dom = g["W4"]["P4"] / g["W4"]["P3"]
    c.ok(g["W4"]["P4"] >= 1.2 * g["W4"]["P3"], f"dominance {dom:.2f}x: F4 decisive edge clears 1.2x F3's carried loss")
    c.ok(g["call"] == "F4" and g["runner_up"] == "F3", "call F4, runner-up F3")
    gap_r = rnd(g["W4"]["P4"], "orders") - rnd(g["W4"]["P3"], "orders")
    c.ok(gap_r == rnd(g["gap"], "orders"), f"rounded gap {gap_r} equals the rounded unrounded gap {g['gap']:.3f}")
    for k, v in (("call W4", g["W4"]["P4"]), ("runner-up W4", g["W4"]["P3"]), ("gap", g["gap"])):
        c.ok(dist(v, "orders") >= MID["orders"], f"{k} {v:.3f} mid-bin")
    # the killing fact of each rung
    c.ok(g["W4"]["P1"] < 0.2 * R["R0"]["F1"] / 4, "R0 killer: by W4 the address check costs a fraction of its stage-table weekly excess")
    c.ok(abs(g["W4"]["P5"]) < 0.25 * R["R2"]["F5"], "R2 killer: mixed baskets against their own baseline lose a fraction of their visitor-type gap")
    w4c = s[(s.week == "W4") & s.club]
    share = float(w4c.P4.mean())
    c.ok(0.5 <= share <= 0.75, f"R3 killer: {share:.1%} of W4 club sessions resolve to store accounts")
    rec["p4_share_w4"] = share
    rec["dominance"] = dom
    return R, rec


def grid(c, ctx, g, s, R):
    out = {}
    st_noclub = N.stage_table(s, exclude_club=True)
    cells = {
        (0, 0, 0): N.rank(R["R0"])[0],
        (1, 0, 0): N.rank(R["R1"])[0],
        (0, 1, 0): N.rank(st_noclub)[0],
        (1, 1, 0): N.rank(R["R3"])[0],
        (0, 0, 1): N.rank(N.stage_table(s))[0],
        (1, 1, 1): N.rank(R["R4"])[0],
    }
    w4 = s[s.week == "W4"]
    p2_store = float(w4.P2.sum() * g["base"]["store"] - w4[w4.P2].conv.sum())
    d101 = dict(R["R4"])
    d101["F2"] = p2_store
    cells[(1, 0, 1)] = N.rank(d101)[0]
    four = {N.FIX[p]: g["L1"][p] for p in N.POPS}
    cells[(0, 1, 1)] = N.rank(four)[0]
    want = {(0, 0, 0): "F1", (1, 0, 0): "F2", (0, 1, 0): "F1", (1, 1, 0): "F3", (0, 0, 1): "F1",
            (1, 0, 1): "F4", (1, 1, 1): "F4"}
    rule = {(0, 0, 0): "the flag cohorts drain (R0)", (1, 0, 0): "the close-outs' visitor-type booking (R1)",
            (0, 1, 0): "the sizing line's latest complete week", (1, 1, 0): "the token chain (R3)",
            (0, 0, 1): "the shortlist's population lines", (1, 0, 1): "right name, F2 cells break the close-outs",
            (1, 1, 1): "the answer"}
    for k, v in want.items():
        c.ok(cells[k] == v, f"grid cell D{k[0]} V{k[1]} T{k[2]} names {v}: {rule[k]}")
    c.ok(cells[(0, 1, 1)] in ("F4", "F1"), f"grid cell D0 V1 T1 (four-week window) names {cells[(0, 1, 1)]}")
    out["cells"] = {f"D{k[0]}V{k[1]}T{k[2]}": v for k, v in cells.items()}
    out["four_week"] = {k: round(v, 2) for k, v in four.items()}
    # partial applications
    bs, share = N.by_step_cell(s, g["base"], g["L"])
    c.ok(N.rank(bs)[0] != "F4", f"by-step booking of re-attached losses names {N.rank(bs)[0]}, not F4")
    out["by_step"] = {k: round(v, 2) for k, v in bs.items()}
    sig = signatures_cell(ctx, s, g)
    c.ok(sig["F4"] < g["W4"]["P3"], f"signatures-only re-attachment finds F4 {sig['F4']:.1f} < F3 {g['W4']['P3']:.1f}")
    out["signatures_only_F4"] = round(sig["F4"], 2)
    p4_store = float(w4.P4.sum() * g["base"]["store"] - w4[w4.P4].conv.sum())
    c.ok(p4_store > g["W4"]["P3"] and not same(p4_store, g["W4"]["P4"], "orders"),
         f"P4 against the store rate {p4_store:.1f} still names F4 with a wrong figure")
    out["p4_store_rate"] = round(p4_store, 2)
    return out


def signatures_cell(ctx, s, g):
    pk = ctx["pk"]
    ab = pk["address"]
    fp2a = dict(zip(ab.address_fp, ab.account_id))
    cards = pk["cards"]
    cfp2a = dict(zip(cards.card_fp, cards.account_id))
    au = pk["auth"].drop_duplicates("attempt_ref")
    s_card = au.groupby("session_id").card_fp.first().to_dict()
    club = s[s.club].copy()
    acc = []
    for sid, f in zip(club.session_id, club.address_fp):
        a = fp2a.get(f) if isinstance(f, str) else None
        if a is None and sid in s_card:
            a = cfp2a.get(s_card[sid])
        acc.append(a)
    club["sig"] = acc
    w4 = club[(club.week == "W4") & club.sig.notna()]
    accs = set(club.sig.dropna())
    b = s[s.week.isin(P.BASE) & s.signed_in & s.account_id.isin(accs)]
    base = b.conv.mean()
    return {"F4": float(len(w4) * base - w4.conv.sum()), "club": club}


# --------------------------------------------------------------------------- calibration organs

def calibration(c, ctx, g, s):
    rec = {}
    for k, b in ctx["books"].items():
        c.ok(abs(b["visitor_type"]) < 0.5, f"close-out {k}: visitor-type booking reproduces the booked 0 ({b['visitor_type']:+.2f})")
        c.ok(b["store_rate"] >= 60, f"close-out {k}: store-rate booking misses by {b['store_rate']:+.1f}")
        c.ok(b["all_new"] <= -30, f"close-out {k}: all-new booking misses by {b['all_new']:+.1f}")
        c.ok(abs(b["signed_in_share_returning"] - b["store_signed_in_share_returning"]) < 2.0,
             f"close-out {k} blind: partner returning visitors signed in at {b['signed_in_share_returning']:.1f}% "
             f"against the store's {b['store_signed_in_share_returning']:.1f}%")
        rec[k] = {kk: round(v, 3) if isinstance(v, float) else v for kk, v in b.items()}
    # resemblance: the club app looks like both partners on every visible column, so lookup books it as dilution
    w4c = s[(s.week == "W4") & s.club]
    conv = 100 * w4c.conv.mean()
    pconv = [b["partner_conversion"] for b in ctx["books"].values()]
    c.ok(w4c.new_visitor.mean() > 0.5 and conv < 100 * g["base"]["store"] and min(pconv) > conv * 0.5,
         f"club app resembles the partners: new-visitor majority, blended conversion {conv:.2f}%")
    lookup_f2 = float(len(w4c) * g["base"]["NV"] - w4c.conv.sum())
    c.ok(lookup_f2 < g["W4"]["P3"], f"lookup transfer books the club as dilution ({lookup_f2:.1f} orders): the decoy reading")
    rec["club_conversion_w4"] = round(conv, 3)
    for k, e in ctx["effs"].items():
        miss = abs(e["calendar"] - e["cohort"]) / abs(e["cohort"])
        c.ok(miss >= 0.3, f"release {k}: calendar-date measurement misses the cohort measurement by {miss:.0%}")
        rec[k] = {kk: round(v, 4) for kk, v in e.items()}
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(ctx["tgt"], ctx["F"]["release_log"]), read_only=True)
    filed = {r[0]: r[6] for r in wb["Releases"].iter_rows(min_row=2, values_only=True)}
    for k, e in ctx["effs"].items():
        c.ok(abs(filed[k] - e["cohort"]) <= 0.02 * abs(e["cohort"]), f"release {k}: filed effect {filed[k]} reproduced by cohort exposure")
    rec["twin"] = twin_pair(c, ctx, s, g)
    rec["drain"] = drain(c, s, g)
    return rec


def cp4_at(pk, when):
    ab = pk["address"]
    ab = ab[ab.is_default].copy()
    ab["tl"] = pd.to_datetime(ab.changed_at) + pd.Timedelta(hours=1)
    ab = ab[ab.tl <= when].sort_values("tl").groupby("account_id").postcode.last()
    return set(ab[ab.str.fullmatch(r"\d{4}")].index)


def twin_pair(c, ctx, s, g):
    pk = ctx["pk"]
    fl = pk["flags"]
    live = {d["cohort"]: pd.Timestamp(d["enabled_at"][:19]) for d in fl["rollout"]}
    c.ok(live[6] == live[8], "twin cohorts 6 and 8 share their exposure time")
    b = s[s.week.isin(P.BASE) & s.signed_in]
    mix = lambda col, ch: b[b.cohort == ch][col].value_counts(normalize=True)
    dd = (mix("device", 6) - mix("device", 8)).abs().max()
    ds = (mix("traffic_source", 6) - mix("traffic_source", 8)).abs().max()
    c.ok(dd < 0.03 and ds < 0.03, f"twin cohorts match on device mix ({dd:.3f}) and source mix ({ds:.3f})")
    cp4 = cp4_at(pk, live[6])
    cur = {d["account_id"]: d["cohort"] for d in fl["assignments"]}
    mv = {d["account_id"]: d["from_cohort"] for d in fl["assignment_moves"]}
    coh = {a: mv.get(a, cc) for a, cc in cur.items()}
    n6 = sum(1 for a in cp4 if coh.get(a) == 6)
    n8 = sum(1 for a in cp4 if coh.get(a) == 8)
    loss = {}
    for ch in (6, 8):
        x = s[s.signed_in & (s.cohort == ch) & s.week.isin(["W1", "W2"])]
        loss[ch] = (float(len(x) * g["base"]["SI"] - x.conv.sum()), len(x))
    r = loss[6][0] / loss[8][0]
    c.ok(n6 / n8 >= 2.0 and r >= 2.0, f"twin cohorts: CP4 accounts {n6} vs {n8}, losses over their first two live weeks {loss[6][0]:.1f} vs {loss[8][0]:.1f}")
    pred8 = loss[6][0] / loss[6][1] * loss[8][1]
    miss = abs(pred8 - loss[8][0]) / loss[8][0]
    c.ok(miss >= 0.4, f"lookup transfer from cohort 6 misses cohort 8 by {miss:.0%}")
    return {"cp4": [n6, n8], "loss": [round(loss[6][0], 2), round(loss[8][0], 2)], "lookup_miss": round(miss, 3)}


def drain(c, s, g):
    fl = None
    x = s[s.signed_in & s.week.isin(P.REVIEW) & s.live].copy()
    live = {1: pd.Timestamp(P.COHORT_LIVE[1]), 5: pd.Timestamp(P.COHORT_LIVE[5]), 9: pd.Timestamp(P.COHORT_LIVE[9])}
    t0 = x.cohort.map(lambda ch: pd.Timestamp(P.COHORT_LIVE[ch]))
    x["wk"] = ((x.t - t0).dt.days // 7).clip(upper=3)
    per = x.groupby("wk").apply(lambda d: g["base"]["SI"] - d.conv.mean(), include_groups=False)
    c.ok(per[2] < 0.3 * per[0], f"event study by weeks since exposure: loss per session {per[0]*100:.2f} pts in week 1, "
         f"{per[2]*100:.2f} in week 3")
    return {int(k): round(float(v) * 100, 3) for k, v in per.items()}


# --------------------------------------------------------------------------- convergence and identity

def convergence(c, ctx, g, s):
    pk = ctx["pk"]
    rec = {}
    si = g["base"]["SI"]
    for p in ("P1", "P3", "P4"):
        c.ok(abs(g["base"][p] - si) * 100 < 0.01, f"{p} own baseline {g['base'][p]*100:.3f}% equals the signed-in rate to 0.01 pts")
    alts = {"signed_in": N.figures(pk, s, base_mode="signed_in"),
            "three_weeks": N.figures(pk, s, base_weeks=("B2", "B3", "B4")),
            "five_weeks": N.figures(pk, s, base_weeks=("B0", "B1", "B2", "B3", "B4"))}
    for name, a in alts.items():
        bad = [(k, p) for k in ("W4", "L1", "L2") for p in N.POPS if not same(a[k][p], g[k][p], "orders")]
        if name == "signed_in":
            c.ok(not bad and a["call"] == "F4", f"{name} baseline leaves every lost-order figure in its bin {bad}")
        else:
            c.ok(a["call"] == "F4" and a["runner_up"] == "F3" and a["W4"]["P4"] >= 1.2 * a["W4"]["P3"],
                 f"{name} baseline still names F4 over F3; moved cells {bad} break the sizing line's four weeks")
        rec[name] = {"W4": {p: round(a["W4"][p], 3) for p in N.POPS}, "moved": [list(x) for x in bad]}
    b = s[s.week.isin(P.BASE)]
    meanw = {}
    for p in ("P1", "P3", "P4"):
        acc = N.pop_accounts(s, p)
        meanw[p] = np.mean([b[(b.week == w) & b.signed_in & b.account_id.isin(acc)].conv.mean() for w in P.BASE])
    c.ok(all(abs(meanw[p] - g["base"][p]) * 100 < 0.05 for p in meanw), "pooled and mean-of-weeks baselines agree")
    for p in ("P1", "P3", "P4"):
        w4 = s[(s.week == "W4") & s[p]]
        alt = len(w4) * meanw[p] - w4.conv.sum()
        c.ok(same(alt, g["W4"][p], "orders"), f"{p} W4 in its bin under the mean-of-weeks baseline")
    # the week's own customers: each week's population accounts against their own August
    for p in ("P1", "P3", "P4"):
        col = "chain_account" if p == "P4" else "account_id"
        bad = []
        for w in P.REVIEW:
            x = s[(s.week == w) & s[p]]
            bw = b[b.signed_in & b.account_id.isin(set(x[col].dropna()))].conv.mean()
            alt = len(x) * bw - x.conv.sum()
            if not same(alt, g["L"][p][w], "orders"):
                bad.append((w, round(alt, 3), round(g["L"][p][w], 3)))
            rec.setdefault("week_sets", {})[f"{p} {w}"] = round(alt, 3)
        c.ok(not bad, f"{p}: each week's own customers as baseline leave every weekly figure in its bin {bad}")
    # unit: per account-week for P4
    acc = N.pop_accounts(s, "P4")
    bo = b[b.signed_in & b.account_id.isin(acc)].conv.sum() / 4
    w4 = s[s.week == "W4"]
    wo = w4[(w4.signed_in & w4.account_id.isin(acc)) | (w4.P4 & w4.chain_account.isin(acc))].conv.sum()
    c.ok((bo - wo) >= 1.2 * g["W4"]["P3"] and not same(bo - wo, g["W4"]["P4"], "orders"),
         f"P4 per account-week {bo - wo:.2f} still names F4 with a figure that breaks the guide's per-session conversion")
    rec["per_account_week"] = round(bo - wo, 3)
    # identity closures
    t_sep = set(pk["tokens_sep"].token)
    t_aug = set(pk["tokens_aug"].token)
    w4c = s[(s.week == "W4") & s.club]
    c.ok(w4c.token.isin(t_sep).all(), "every W4 club token is in the September token report")
    c.ok(set(pk["tokens_aug"].issued_at.str[:10]) == {"2026-08-31"}, "the August token report covers 31 August only")
    pr = pk["profiles"].dropna(subset=["club_member_no"])
    c.ok(pr.club_member_no.is_unique, "every member number sits on at most one loyalty profile")
    T = ctx["T"]
    mem = T["acc"][T["acc"].klass == "MEMBER"].account_id
    c.ok(mem.isin(pr.account_id).all(), "every member with a store account carries the member number on the profile")
    sig = signatures_cell(ctx, s, g)["club"]
    obs = sig[sig.sig.notna() | sig.address_fp.notna()]
    dis = int(((obs.sig != obs.chain_account) & (obs.sig.notna() | obs.chain_account.notna())
               & ~(obs.sig.isna() & obs.P4 & obs.address_fp.isna())).sum())
    p2obs = int((obs.P2 & obs.sig.notna()).sum())
    p4obs_bad = int((obs.P4 & obs.address_fp.notna() & (obs.sig != obs.chain_account)).sum())
    c.ok(p2obs == 0 and p4obs_bad == 0, f"signatures agree with the token chain on every observable session ({len(obs)})")
    return rec


def clean_and_lens(c, ctx, g, s, R):
    # clean-data test: the session log repaired with the account behind every club session
    rep = s.copy()
    rep.loc[rep.P4, "account_id"] = rep.loc[rep.P4, "chain_account"]
    naive_rep = N.rank(N.stage_table(rep))[0]
    acc = set(rep.account_id[rep.P4])
    w4 = rep[(rep.week == "W4") & rep.P4]
    b = rep[rep.week.isin(P.BASE) & rep.signed_in & rep.account_id.isin(acc)].conv.mean()
    ans = dict(R["R4"])
    ans["F4"] = float(len(w4) * b - w4.conv.sum())
    ans_rep = N.rank(ans)[0]
    c.ok(ans_rep == "F4" == g["call"] and naive_rep == "F1" == N.rank(R["R0"])[0] and ans_rep != naive_rep,
         "clean-data test on the session log: answer(repaired) = answer = F4, naive(repaired) = naive = F1")
    # lens swap: P4 is a population reached through a join, not the club sessions under another lens
    club = s[(s.week == "W4") & s.club]
    c.ok(club.P4.any() and (~club.P4).any(), "P4 is a strict subset of the club sessions")
    best = 0.0
    for col in ("device", "traffic_source", "new_visitor", "signed_in", "furthest_step", "basket_skus"):
        ct = pd.crosstab(club[col], club.P4)
        best = max(best, float(ct.max(axis=1).sum() / len(club)))
    c.ok(best < 0.70, f"no single session-log column separates P4 (best accuracy {best:.2f})")
    gen = s[s.club & ~s.bot]
    c.ok(int(gen.P5.sum()) == 0, "no genuine club session holds a pre-order line")
    return {"lens_best_accuracy": round(best, 3)}


def separation(c, ctx, g, s, variants):
    pk = ctx["pk"]
    main_w = set(P.BASE) | {"W4"}
    allw = N.enrich(pk, N.opts(keep_bots=True))
    c.ok(int((allw.bot & allw.week.isin(main_w)).sum()) == 0, "separation: zero edge-flagged sessions in the baseline weeks and W4")
    ab = pk["address"]
    ab = ab[ab.event == "updated"].copy()
    ab["R"] = pd.to_datetime(ab.changed_at) + pd.Timedelta(hours=1)
    cp4acc = set(ctx["T"]["acc"].account_id[ctx["T"]["acc"].klass == "CP4"])
    rs = ab[ab.account_id.isin(cp4acc)]
    ms = s[s.week.isin(main_w) & s.signed_in][["account_id", "t"]].merge(rs[["account_id", "R"]], on="account_id")
    hits = int(((ms.R > ms.t) & (ms.R <= ms.t + pd.Timedelta(hours=2))).sum())
    c.ok(hits == 0, "separation: zero address re-saves within two hours after a baseline-week or W4 session")
    ce = pk["card_changes"]
    ce = ce[ce.event == "set_default"]
    wk = N.weeks_of(pd.to_datetime(ce["at"]))
    c.ok(int(pd.Series(wk).isin(list(main_w)).sum()) == 0, "separation: zero default-card changes in the baseline weeks and W4")
    for name, o in variants.items():
        vs = o["s"]
        cols = ["P1", "P2", "P3", "P4", "P5"]
        a = s[s.week == "W4"].set_index("session_id")[cols]
        bb = vs[vs.week == "W4"].set_index("session_id")[cols].reindex(a.index)
        diff = int((a != bb).any(axis=1).sum())
        vl = N.losses(vs, N.baselines(vs), weeks=["W4"])
        off = [p for p in N.POPS if not same(vl[p]["W4"], g["W4"][p], "orders")]
        c.ok(diff == 0 and not off, f"separation: the {name} reading changes no W4 classification and no W4 lost-order "
             f"figure ({diff}, {off})")
    return {}


# --------------------------------------------------------------------------- the ask layer

DEVICES = {  # device: (enrich options, finance flags, auth flags, fold club, asks it moves)
    "DV1 automated sessions kept": (dict(keep_bots=True), (), (), False, {"A1", "A4"}),
    "DV2 address clock read raw": (dict(clock_raw=True), (), (), False, {"A1"}),
    "DV3 saved cards read as current": (dict(cards_current=True), (), (), False, {"A1"}),
    "DV4 shipment rows summed": ({}, ("rows_summed",), (), False, {"A2"}),
    "DV5 gross rows kept": ({}, ("gross_kept",), (), False, {"A2"}),
    "DV6 decoupled challenges dropped": ({}, (), ("C_only",), False, {"A3"}),
    "DV7 re-sent messages counted": ({}, (), ("rows",), False, {"A3"}),
    "DV8 cohorts read as current": (dict(cohort_current=True), (), (), False, {"A4"}),
    "DV9 September token report only": (dict(tokens_sep_only=True), (), (), False, {"A1"}),
    "DV10 catalogue status read as current": (dict(catalogue_current=True), (), (), False, {"A1"}),
    "DV11 first export kept": ({}, ("pre_correction",), (), False, {"A2"}),
    "DV13 club sessions folded into cohorts": ({}, (), (), True, {"A4"}),
}


class Cache:
    def __init__(self, pk, s):
        self.pk = pk
        self.d = {tuple(): s}

    def get(self, o):
        key = tuple(sorted(o.items()))
        if key not in self.d:
            self.d[key] = N.enrich(self.pk, N.opts(**o))
        return self.d[key]


def sheet(pk, s, fin=(), auth=(), fold=False):
    g = N.figures(pk, s, fin_mode=fin or "golden", auth_mode=auth or "golden")
    coh = N.cohort_sheet(s, fold_club=fold, pack=pk)
    return g, ask_figs(g, coh)


def asks(c, ctx, g, cache):
    pk = ctx["pk"]
    base = ask_figs(g, N.cohort_sheet(cache.get({})))
    rec = {"necessity": {}}
    for name, (o, fin, au, fold, mv) in DEVICES.items():
        gd, fd = sheet(pk, cache.get(o), fin, au, fold)
        moved_asks = {k[0] for k in moved(base, fd)}
        c.ok(moved_asks == mv, f"{name}: moves exactly {sorted(mv)} (moved {sorted(moved_asks)})")
        c.ok(gd["call"] == "F4" and same(gd["call_w4"], g["call_w4"], "orders") and same(gd["gap"], g["gap"], "orders"),
             f"{name}: the main call, its figure and the gap are unchanged")
        rec["necessity"][name] = len(moved(base, fd))
    # over-cleaning: drop every session on a touched token or account
    go, fo = sheet(pk, cache.get(dict(drop_touched=True)))
    oc = {k[0] for k in moved(base, fo)}
    c.ok({"A1", "A4"} <= oc, f"over-cleaner lands A1 and A4 outside the golden bins ({sorted(oc)})")
    # composed mishandlings: every pair and the full set, per ask
    a1 = [n for n in DEVICES if "A1" in DEVICES[n][4]]
    for r in (2, len(a1)):
        for combo in itertools.combinations(a1, r):
            o = {}
            for n in combo:
                o.update(DEVICES[n][0])
            _, fc = sheet(pk, cache.get(o))
            c.ok(len(moved(base, fc, "A1")) > 0, f"A1 composed {' + '.join(x.split()[0] for x in combo)} stays off the golden")
    for ask, keys in (("A2", ["DV4 shipment rows summed", "DV5 gross rows kept", "DV11 first export kept"]),
                      ("A3", ["DV6 decoupled challenges dropped", "DV7 re-sent messages counted"])):
        for r in range(2, len(keys) + 1):
            for combo in itertools.combinations(keys, r):
                fin = tuple(f for n in combo for f in DEVICES[n][1])
                au = tuple(f for n in combo for f in DEVICES[n][2])
                _, fc = sheet(pk, cache.get({}), fin, au)
                c.ok(len(moved(base, fc, ask)) > 0, f"{ask} composed {' + '.join(x.split()[0] for x in combo)} stays off the golden")
    for combo in (("DV1", "DV8"), ("DV1", "DV13"), ("DV8", "DV13"), ("DV1", "DV8", "DV13")):
        names = [n for n in DEVICES if n.split()[0] in combo]
        o = {}
        fold = False
        for n in names:
            o.update(DEVICES[n][0])
            fold = fold or DEVICES[n][3]
        _, fc = sheet(pk, cache.get(o), fold=fold)
        c.ok(len(moved(base, fc, "A4")) > 0, f"A4 composed {' + '.join(combo)} stays off the golden")
    # lazy sweep: the natural read of each ask
    lazy = {"A1": (dict(keep_bots=True, clock_raw=True, cards_current=True, tokens_sep_only=True,
                        catalogue_current=True), (), ()),
            "A2": ({}, ("rows_summed", "gross_kept", "pre_correction"), ()),
            "A3": ({}, (), ("rows", "C_only")),
            "A4": (dict(keep_bots=True, cohort_current=True), (), ())}
    rec["lazy"] = {}
    for ask, (o, fin, au) in lazy.items():
        _, fl = sheet(pk, cache.get(o), fin, au)
        mv_ = moved(base, fl, ask)
        n_ask = len([k for k in base if k[0] == ask])
        c.ok(len(mv_) >= (n_ask + 1) // 2 if ask != "A4" else len(mv_) >= 8,
             f"lazy sweep on {ask}: {len(mv_)} of {n_ask} figures off the golden")
        rec["lazy"][ask] = [len(mv_), n_ask]
    # every ask stop distinct from the golden per figure
    stops = {"A2 store-wide order value": ({}, ("store_wide",), ())}
    _, fs = sheet(pk, cache.get({}), ("store_wide",))
    c.ok(len(moved(base, fs, "A2")) >= 1, f"A2 stop: every order valued at the store average lands off the golden ({len(moved(base, fs, 'A2'))} of 5)")
    return rec, base


def pair_simulation(c, ctx, g, cache, base):
    """Cracker and mirror sheets, both swept with the habitual battery: duplicate keys removed
    (shipment rows, re-sent messages, first export kept), unmatched joins chased (both token
    reports), nothing beyond that."""
    pk = ctx["pk"]
    o = dict(keep_bots=True, clock_raw=True, cards_current=True, catalogue_current=True, cohort_current=True)
    _, crack = sheet(pk, cache.get(o), ("gross_kept", "pre_correction"), ("C_only",))
    s_m = cache.get(o)
    gm = N.figures(pk, s_m, fin_mode=("gross_kept", "pre_correction"), auth_mode=("C_only",))
    mirror = dict(crack)
    for k in list(mirror):
        if k[0] in ("A1", "A2", "A3") and k[2] in ("P2", "P4"):
            mirror[k] = (-999.0, mirror[k][1])
    n = len(base)
    lc = sum(1 for k in base if same(base[k][0], crack[k][0], base[k][1])) / n
    lm = sum(1 for k in base if same(base[k][0], mirror[k][0], base[k][1])) / n
    chart_free = 5
    lc_w = (lc * n + chart_free) / (n + chart_free + 2)
    lm_w = (lm * n + 2) / (n + chart_free + 2)
    cracker = 38 + 7 + 55 * lc_w
    mirror_s = 5 + 7 + 55 * lm_w
    pair = (cracker + mirror_s) / 2
    c.ok(pair <= 45, f"pair simulation: cracker {cracker:.1f}, mirror {mirror_s:.1f}, pair {pair:.1f}")
    return {"cracker": round(cracker, 1), "mirror": round(mirror_s, 1), "pair": round(pair, 1),
            "leak_cracker": round(lc, 3), "leak_mirror": round(lm, 3)}


def battery_and_markers(c, ctx, cache):
    pk = ctx["pk"]
    s = pk["sessions"]
    c.ok(s.session_id.is_unique, "battery: session keys unique")
    sb = cache.get(dict(keep_bots=True))
    days = sb[sb.t.dt.date.isin([dt.date(2026, 8, 31), dt.date(2026, 9, 1), dt.date(2026, 9, 16), dt.date(2026, 9, 17),
                                 dt.date(2026, 9, 18)])]
    worst = 0.0
    for col in ("device", "traffic_source", "signed_in", "new_visitor", "furthest_step"):
        sh = days.groupby(col).bot.mean().max()
        worst = max(worst, float(sh))
    c.ok(worst < 0.5, f"no marker: no value of a session-log column isolates the automated sessions ({worst:.2f})")
    # referee: the provider's weekly totals tie only on the attempt grain with D counted
    import pypdf
    txt = "".join(pg.extract_text() for pg in pypdf.PdfReader(os.path.join(ctx["tgt"], ctx["F"]["psp"])).pages)
    m = re.search(r"Total\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)\s+([\d,]+)", txt)
    tot = [int(x.replace(",", "")) for x in m.groups()]
    a = pk["auth"]
    d = a.drop_duplicates("attempt_ref")
    c.ok(tot[0] == len(d) and tot[2] == int(d.three_ds_status.isin(["C", "D"]).sum()),
         "referee: the provider's totals tie with attempts counted once and decoupled challenges counted")
    c.ok(tot[0] != len(a) and tot[2] != int(d.three_ds_status.eq("C").sum()),
         "referee: rows, or challenges without D, do not tie")
    return {"psp_totals": tot}


# --------------------------------------------------------------------------- the pack itself

def text_of(path):
    ext = path.rsplit(".", 1)[1]
    if ext in ("csv", "md", "txt", "json"):
        return open(path, encoding="utf-8").read()
    if ext in ("docx", "xlsx"):
        with zipfile.ZipFile(path) as z:
            return " ".join(re.sub(r"<[^>]+>", " ", z.read(n).decode("utf-8", "ignore"))
                            for n in z.namelist() if n.endswith(".xml"))
    if ext == "pdf":
        import pypdf
        return " ".join(pg.extract_text() for pg in pypdf.PdfReader(path).pages)
    if ext == "parquet":
        import pyarrow.parquet as pq
        return " ".join(pq.read_schema(path).names)
    return ""


def pack_gates(c, ctx):
    tgt = ctx["tgt"]
    F = ctx["F"]
    meta = ctx["meta"]
    files = sorted(os.listdir(tgt))
    fmts = {f.rsplit(".", 1)[1] for f in files}
    c.ok(len(files) >= 10, f"input gate: {len(files)} files")
    c.ok(len(fmts) >= 3, f"input gate: {len(fmts)} formats {sorted(fmts)}")
    import pyarrow.parquet as pq
    nrows = pq.ParquetFile(os.path.join(tgt, F["sessions"])).metadata.num_rows
    c.ok(nrows >= 25000, f"input gate: the session log carries {nrows:,} rows")
    ds = meta["distractor_files"]
    c.ok(len(ds) >= 2 and all(d in files for d in ds), "input gate: two distractors named in metadata.json, both shipped")
    texts = {f: text_of(os.path.join(tgt, f)) for f in files}
    c.ok(not any("distractor" in (f + t).lower() for f, t in texts.items()), "no file name or text calls a file a distractor")
    c.ok(set(ds).isdisjoint(set(ctx["pk"].keys())), "distractors are never read by the golden computation")
    c.ok(meta["deliverables"] == ["q4_sprint_call.html", "checkout_fall_workings.xlsx"], "metadata names the two deliverables")
    # single-statement invariant: each load-bearing fact in exactly one file
    facts = {"sizing line": "costing us the most orders now",
             "new-fan dilution basis": "Partner traffic is new-fan dilution",
             "automated exclusion": "flags as automated are excluded",
             "address book in UTC": "changed_at is UTC",
             "decoupled challenge code": "approved the payment in the issuer's banking app",
             "latest export of record": "latest export of the order is the one of record",
             "net from the release": "net of VAT on rows exported from the 10 August",
             "cohort as of the session": "latest assignment at or before the session",
             "attempt key": "attempt_ref is one payment",
             "token to member number": "with the member number it was issued to"}
    for k, ph in facts.items():
        hits = [f for f, t in texts.items() if ph.lower() in re.sub(r"\s+", " ", t).lower()]
        c.ok(len(hits) == 1, f"single statement: '{k}' in exactly one file ({hits})")
    banned = ["in-app browser", "signed out", "own browser", "existing customer", "existing account",
              "logged out", "resolves to an account", "\u2014"]
    for ph in banned:
        hits = [f for f, t in texts.items() if ph in t.lower()]
        c.ok(not hits, f"anti-signpost: '{ph}' appears in no shipped file ({hits})")
    fixes = ["postcode lookup", "landing page", "sdk", "one-time-code", "split dispatch"]
    rankers = [f for f, t in texts.items() if sum(x in t.lower() for x in fixes) >= 3 and f != F["shortlist"]]
    c.ok(not rankers, f"no shipped artifact other than the shortlist names three or more fixes ({rankers})")
    sl = texts[F["shortlist"]]
    c.ok(not re.search(r"\b\d{2,}\s*(orders|%|eur)", sl.lower()), "the shortlist carries no figure that ranks the fixes")
    # container hygiene: no writer signature, every embedded date inside the fiction's band
    sigs = ["openpyxl", "python-docx", "reportlab", "matplotlib", "parquet-cpp", "pandas", "arrow", "xlsxwriter"]
    for f in files:
        raw = open(os.path.join(tgt, f), "rb").read()
        if f.endswith((".docx", ".xlsx")):
            with zipfile.ZipFile(os.path.join(tgt, f)) as z:
                raw = b"".join(z.read(n) for n in z.namelist() if n.startswith("docProps/"))
        if f.endswith((".csv", ".md", ".txt", ".json")):
            continue
        low = raw.lower()
        c.ok(not any(x.encode() in low for x in sigs), f"container of {f} names no writer")
        ds_ = [m.decode()[:10] for m in re.findall(rb"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", raw)]
        ds_ += [f"{m.decode()[2:6]}-{m.decode()[6:8]}-{m.decode()[8:10]}" for m in re.findall(rb"D:\d{14}", raw)]
        c.ok(all("2025-06-01" <= d <= "2026-09-30" for d in ds_), f"container dates of {f} inside the fiction {sorted(set(ds_))}")
    sc = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))),
                      ".claude", "skills", "reduce-house-fixes", "scripts", "scrub_producer_metadata.py")
    r = subprocess.run([sys.executable, sc, tgt], capture_output=True, text=True)
    c.ok(r.returncode == 0 and not re.search(r"(?i)openpyxl|python-docx|reportlab", r.stdout.split("producer")[-1] if False else
                                              "\n".join(l for l in r.stdout.splitlines() if "signature" in l.lower() and "none" not in l.lower())),
         "scrub_producer_metadata audit finds no producer signature")
    # the register knows every file (H9) and no record runs past the export (H16)
    reg = pd.read_csv(os.path.join(tgt, F["register"]))
    c.ok(set(reg.file) == set(files) - {F["register"]}, "the extract register lists exactly the shipped files")
    late = []
    for f in files:
        if f.endswith(".csv") or f.endswith(".md") or f.endswith(".txt") or f.endswith(".json"):
            ds_ = re.findall(r"\b(20\d\d-\d\d-\d\d)", texts[f])
            if ds_ and max(ds_) > "2026-09-30" and f not in (F["stock"], F["release_log"]):
                late.append((f, max(ds_)))
    c.ok(not late, f"no shipped extract carries a date after the export ({late})")
    s = pd.read_parquet(os.path.join(tgt, F["sessions"]))
    c.ok(s.started_at.max() < "2026-09-28" and s.started_at.min() >= "2026-07-27", "the session log spans 27 July to 27 September")
    # the context artifact reproduces from the spine
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(tgt, F["dashboard"]), read_only=True)
    rows = [r for r in wb["Weekly"].iter_rows(min_row=6, values_only=True) if r[0]]
    wk = ctx["dash"][0]
    c.ok([tuple(r[:4]) for r in rows] == [(w["week"], w["sessions"], w["orders"], w["conversion"]) for w in wk],
         "the trading dashboard's weekly table reproduces from the session log")
    lab = " ".join(str(x) for x in next(wb["Weekly"].iter_rows(min_row=3, max_row=3, values_only=True)))
    c.ok("does not size causes" in lab, "the dashboard is labelled in-file for the question it answers")
    counts = [x for x in (len(ctx["pk"][k]) for k in ("edge", "tokens_aug", "tokens_sep", "profiles", "address", "cards",
                                                        "card_changes", "auth", "finance"))]
    c.ok(len(set(counts)) == len(counts), "no two extracts share a row count")
    return {"files": len(files), "formats": sorted(fmts), "rows": int(nrows)}


def mid_bins(c, g, coh):
    figs = ask_figs(g, coh)
    for k, (v, kind) in figs.items():
        if k[1] == "n":
            continue
        c.ok(dist(v, kind) >= MID[kind], f"{'/'.join(map(str, k))} = {v:.4f} sits mid-bin")
    for p in N.POPS:
        wk = [g["L"][p][w] for w in P.REVIEW]
        c.ok(sum(rnd(x, "orders") for x in wk) == rnd(g["L1"][p], "orders"),
             f"{p} four-week loss: rounded-then-summed equals summed-then-rounded")
    sh = sorted(g["L4"].values())
    c.ok(all(b - a >= 0.3 for a, b in zip(sh, sh[1:])), "challenge shares at least 0.3 points apart")
    c.ok(all(abs(v * 10 - round(v * 10)) > 1e-6 for v in g["L4"].values()), "no challenge share is an exact round figure")


def run(ctx):
    c = Chk()
    s = ctx["s"]
    pk = ctx["pk"]
    g = N.figures(pk, s)
    cache = Cache(pk, s)
    R, rec_l = ladder(c, ctx, g, s)
    rec_g = grid(c, ctx, g, s, R)
    rec_c = calibration(c, ctx, g, s)
    rec_v = convergence(c, ctx, g, s)
    rec_x = clean_and_lens(c, ctx, g, s, R)
    variants = {n: {"s": cache.get(o)} for n, o in (("raw clock", dict(clock_raw=True)),
                                                     ("current saved card", dict(cards_current=True)),
                                                     ("September-only token", dict(tokens_sep_only=True)),
                                                     ("current catalogue", dict(catalogue_current=True)),
                                                     ("current cohort", dict(cohort_current=True)))}
    separation(c, ctx, g, s, variants)
    rec_a, base = asks(c, ctx, g, cache)
    rec_p = pair_simulation(c, ctx, g, cache, base)
    rec_b = battery_and_markers(c, ctx, cache)
    coh = N.cohort_sheet(s)
    mid_bins(c, g, coh)
    rec_k = pack_gates(c, ctx)
    c.ok(c.n >= 40, f"{c.n} assertions")
    chart = {N.FIX[p]: {w: round(g["L"][p][w], 3) for w in P.WEEK_NAMES[1:]} for p in N.POPS}
    return {
        "assertions": c.n,
        "call": g["call"], "call_w4": g["W4"]["P4"], "runner_up": g["runner_up"], "runner_up_w4": g["W4"]["P3"],
        "gap": g["gap"],
        "L1": g["L1"], "L2": g["L2"], "L3": g["L3"], "L4": g["L4"], "L4_kn": g["L4_kn"], "aov": g["aov"],
        "W4": g["W4"], "base": {k: v for k, v in g["base"].items() if not isinstance(v, dict)},
        "cohorts": {str(k): [n, v] for k, (n, v) in coh.items()},
        "chart": chart, "ladder": rec_l, "grid": rec_g, "calibration": rec_c, "convergence": rec_v,
        "clean_lens": rec_x, "asks": rec_a, "pair": rec_p, "referee": rec_b, "pack": rec_k,
        "tune_log": ctx["tlog"], "log": c.log,
    }
