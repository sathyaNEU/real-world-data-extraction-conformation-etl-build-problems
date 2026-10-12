"""The three deliverables: the staff determination for the hearing board, the workbook behind it, and the chart of
the season's odor events by wind direction."""
import datetime as dt
import io
import math
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import analysis as AN
import core
import docs as D
import logic as L
import postprocess as PP

DOCX = "odor_source_determination_2026.docx"
XLSX = "odor_events_2026.xlsx"
PNG = "odor_events_by_wind_direction_2026.png"
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September"]
NAME = {k: core.FAC[k]["short"] for k in core.FKEYS}
LONG = {k: core.FAC[k]["name"] for k in core.FKEYS}


def g(x):
    return "{:,}".format(int(x))


def poss(name):
    return name + ("'" if name.endswith("s") else "'s")


def chart_png(F):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    counts = F["sectors"]
    theta = [math.radians(22.5 * i) for i in range(16)]
    fig = plt.figure(figsize=(6.6, 6.9), dpi=150)
    ax = fig.add_subplot(111, projection="polar")
    ax.set_theta_zero_location("N")
    ax.set_theta_direction(-1)
    call = F["call"]
    hot = {i for i in range(16) if core.angdiff(22.5 * i, core.FAC[call]["bearing"]) <= core.UPWIND_DEG}
    cols = ["#C9783A" if i in hot else "#8A9BA8" for i in range(16)]
    ax.bar(theta, counts, width=math.radians(20.5), color=cols, edgecolor="white", linewidth=0.6, zorder=2)
    top = max(counts) * 1.18
    ax.set_ylim(0, top)
    ax.set_xticks(theta)
    ax.set_xticklabels(AN.SECTORS, fontsize=8)
    ax.tick_params(axis="y", labelsize=6.5, colors="#666666")
    ax.set_rlabel_position(292.5)
    for i, v in enumerate(counts):
        if v:
            ax.text(theta[i], v + top * 0.035, str(v), ha="center", va="center", fontsize=6.6, color="#333333")
    for k in core.FKEYS:
        b = math.radians(core.FAC[k]["bearing"])
        strong = k == call
        ax.plot([b, b], [0, top * 0.97], color="#9C4A12" if strong else "#4A5560", lw=1.6 if strong else 0.8,
                ls="-" if strong else (0, (3, 2)), zorder=3)
        ax.text(b, top * 1.06, NAME[k], ha="center", va="center", fontsize=7.2,
                fontweight="bold" if strong else "normal", color="#9C4A12" if strong else "#4A5560")
    fig.suptitle("Ridgeway Organics is behind the 2026 odor complaints in Harlow Falls", fontsize=11, x=0.03, ha="left",
                 y=0.985, fontweight="bold")
    fig.text(0.03, 0.935, "Odor events, May 1 to September 30, 2026, by the direction the wind blew from at the Harlow "
                          "Falls station (%s events).\nLines mark each facility's bearing from the station; the shaded "
                          "bars are the sectors within 30 degrees of Ridgeway Organics.\n%d events came in calm hours "
                          "and %d in hours without a usable wind; they are not shown."
             % (g(sum(counts)), F["calm_events"], F["nowind_events"]), fontsize=7.6, ha="left", va="top",
             color="#333333")
    fig.text(0.03, 0.015, "The tallest bars follow the night drainage from the northeast, which blows over the town on "
                          "most nights whether or not anyone complains.\nSource: complaint system export of October 2, "
                          "2026; Harlow Falls station 5-minute records; District odor source list.",
             fontsize=6.8, ha="left", va="bottom", color="#555555")
    fig.subplots_adjust(top=0.83, bottom=0.08)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def _footer(d, text):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    p = d.sections[0].footer.paragraphs[0]
    r = p.add_run(text)
    r.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    run.append(t)
    fld.append(run)
    p._p.append(fld)


def determination(path, F):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    O, E, X, Q = F["O"], F["E"], F["excess"], F["ratio"]
    call, rival = F["call"], F["rival"]
    d = Document()
    stl = d.styles["Normal"]
    stl.font.name = "Calibri"
    stl.font.size = Pt(10.5)
    _footer(d, "Staff determination, odor complaints 2026, page ")
    p = d.add_paragraph()
    r = p.add_run(core.ORG)
    r.bold = True
    r.font.size = Pt(12.5)
    d.add_paragraph("Compliance Division, 1180 Mill Street, Harlow Falls")
    d.add_paragraph("To: Hearing Board\nFrom: Erica Schmitt, Compliance Manager\nPrepared by: Alyssa White, Air Quality "
                    "Specialist\nDate: October 12, 2026")
    p = d.add_paragraph()
    p.add_run("Re: Source of the 2026 odor complaints in Harlow Falls; petition for an order for abatement, %s "
              "(Permit %s)" % (LONG[call], core.FAC[call]["permit"])).bold = True
    d.add_heading("%s is behind the surge in odor complaints; %s comes closest" % (NAME[call], NAME[rival]), level=1)
    d.add_paragraph(
        "Between May 1 and September 30, 2026 the complaints the District took in Harlow Falls form %s odor events "
        "under CP-14. Staff "
        "find that the odor behind the surge in complaints comes from the windrow pad at %s, in the canyon southwest "
        "of town, "
        "which began taking food waste in March, and recommend that the Board order abatement before the 2027 season. "
        "The facility lay upwind of %s odor events, against %s that the wind at the same hours on the season's other "
        "days would have put it upwind of: %s events more than the wind alone explains, %.1f times the expected "
        "number. No other facility comes near. %s comes closest, upwind of %s events against %s expected (%s more, "
        "%.1f times)."
        % (g(F["n_events"]), LONG[call], g(O[call]), "{:.1f}".format(E[call]), "{:.1f}".format(X[call]), Q[call],
           NAME[rival], g(O[rival]), "{:.1f}".format(E[rival]), "{:.1f}".format(X[rival]), Q[rival]))
    d.add_paragraph(
        "%s lay upwind of more events than any other facility (%s), because the drainage wind from the northeast "
        "blows over the town on most nights and residents report most odors in the evening and at night. At those "
        "same hours on the season's other days the wind would have put the landfill upwind of %s events, more than "
        "it was; the landfill's count measures the night wind, not its odor. Measured against the wind over all hours "
        "of the season, %s looks strongest, because the north slope's night flow is its only bearing; set against the "
        "same hours, its excess is %s %s. %s count is low for the opposite reason: the up-valley wind puts "
        "it upwind of the town all day, when nobody is home to notice, and the evenings its odor reaches town are the "
        "warm ones when that wind lingers past dark."
        % (NAME["landfill"], g(O["landfill"]), "{:.1f}".format(E["landfill"]), NAME["render"],
           "less than half of" if X["render"] < 0.5 * X[call] else "well short of", poss(NAME[call]), poss(NAME[call])))
    tb = d.add_table(rows=1 + len(core.FKEYS), cols=5)
    tb.style = "Light Grid Accent 1"
    hdr = ["Facility", "Odor events it lay upwind of", "Expected from the same hours on other days",
           "Events beyond expected", "Times expected"]
    for j, h in enumerate(hdr):
        tb.cell(0, j).text = h
    order = sorted(core.FKEYS, key=lambda f: -X[f])
    for i, f in enumerate(order, start=1):
        vals = [NAME[f], g(O[f]), "{:.1f}".format(E[f]), "{:.1f}".format(X[f]), "{:.1f}".format(Q[f])]
        for j, v in enumerate(vals):
            tb.cell(i, j).text = v
    for i, row in enumerate(tb.rows):
        for j, cell in enumerate(row.cells):
            for pp in cell.paragraphs:
                if j > 0:
                    pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for run in pp.runs:
                    run.font.size = Pt(9)
                    run.bold = i == 0 or (i == 1)
    sp = d.add_paragraph("Odor events and upwind checks under CP-14 rev. 4: hourly vector-average wind at the Harlow "
                         "Falls station, 30-degree upwind check from each household's geocoded location. Expected: for "
                         "each event, the share of the season's other days on which the facility lay upwind of that "
                         "household at the same clock hour, summed over the events.")
    sp.runs[0].italic = True
    sp.runs[0].font.size = Pt(8.5)
    d.add_heading("The season's record", level=2)
    d.add_paragraph(
        "Of the %s events, %s fell in calm hours and %s in hours without a usable wind (the station's wind monitor "
        "was out from June 9 to 11); CP-14 puts no source upwind in either, and they are left out of every "
        "facility's count. The station recorded a usable wind in %s of the season's hours."
        % (g(F["n_events"]), g(F["calm_events"]), g(F["nowind_events"]), g(sum(F["usable"]))))
    d.add_paragraph(
        "Every facility had a 2026 change, so timing names none of them: the landfill opened Cell 7 in May, the "
        "treatment plant's drying beds came into service in April, the rendering plant added a second cooking shift "
        "in June and the mushroom farm extended its wharf in May. %s food waste contracts with the City and the "
        "County run to June 2030." % poss(NAME[call]))
    fn = d.add_paragraph("CP-14 section 3: an upwind check shows that a source could have reached the household in that "
                         "hour; on its own it does not show which source did. The comparison above is what turns the "
                         "season's upwind checks into a finding.")
    fn.runs[0].font.size = Pt(8.5)
    d.add_paragraph("Attachments: odor_events_2026.xlsx (the season's events, the complaint record by month, the "
                    "station's hours); odor_events_by_wind_direction_2026.png.")
    cp = d.core_properties
    cp.author = "Alyssa White"
    cp.last_modified_by = "Erica Schmitt"
    cp.title = "Staff determination: source of the 2026 odor complaints in Harlow Falls"
    cp.comments = ""
    cp.keywords = ""
    cp.subject = ""
    cp.category = ""
    cp.revision = 4
    cp.created = dt.datetime(2026, 10, 8, 14, 10)
    cp.modified = dt.datetime(2026, 10, 12, 16, 45)
    buf = io.BytesIO()
    d.save(buf)
    D.normalize_docx(buf.getvalue(), path)


def _sheet(ws, head, rows, widths, fmt=None):
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E3E9E4")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    if fmt:
        for col, f in fmt.items():
            for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                for c in row:
                    c.number_format = f
    ws.freeze_panes = "A2"


def workbook(path, F):
    wb = Workbook()
    ws = wb.active
    ws.title = "Complaints by month"
    _sheet(ws, ["Month (2026)", "Complaints about odors noticed that month, inside the District"],
           [[m, n] for m, n in zip(MONTHS, F["months"])] + [["January to September", sum(F["months"])]],
           [22, 30], {2: "#,##0"})
    ws["A11"].font = Font(bold=True)
    ws["B11"].font = Font(bold=True)
    ws2 = wb.create_sheet("Station hours")
    _sheet(ws2, ["Month (2026, local time)", "Hours with a usable wind", "Calm hours"],
           [[m, u, c] for m, u, c in zip(MONTHS[4:], F["usable"], F["calm"])] +
           [["May to September", sum(F["usable"]), sum(F["calm"])]], [24, 16, 12], {2: "#,##0", 3: "#,##0"})
    ws2["A7"].font = Font(bold=True)
    ws3 = wb.create_sheet("Facilities")
    order = sorted(core.FKEYS, key=lambda f: -F["excess"][f])
    _sheet(ws3, ["Facility", "Permit", "Bearing from the station (deg)", "Odor events it lay upwind of",
                 "Expected from the same hours on other days", "Events beyond expected", "Times expected"],
           [[LONG[f], core.FAC[f]["permit"], round(core.FAC[f]["bearing"]), F["O"][f], round(F["E"][f], 1),
             round(F["excess"][f], 1), round(F["ratio"][f], 1)] for f in order],
           [44, 9, 14, 14, 18, 14, 12], {4: "#,##0", 5: "0.0", 6: "0.0", 7: "0.0"})
    ws4 = wb.create_sheet("Odor events")
    rows = []
    for i, e in enumerate(F["events"], start=1):
        dd, hh = L.event_hour(e)
        w = F["hw"].get((dd, hh))
        if w is None:
            state, sp, dr = "No usable wind", None, None
        elif w[0] < core.CALM_MS:
            state, sp, dr = "Calm", round(w[0], 2), None
        else:
            state, sp, dr = "", round(w[0], 2), round(w[1])
        up = [NAME[f] for f in core.FKEYS if L.is_upwind(f, e["lat"], e["lon"], w)]
        rows.append([i, e["household"].title(), e["noticed"], dt.datetime(dd.year, dd.month, dd.day, hh), sp, dr, state,
                     "; ".join(up)])
    _sheet(ws4, ["Event", "Household", "Time noticed (local)", "Event hour (PST)", "Wind speed (m/s)",
                 "Wind direction (deg)", "Calm or no wind", "Facilities upwind"], rows, [7, 30, 18, 17, 10, 10, 15, 48],
           {3: "yyyy-mm-dd hh:mm", 4: "yyyy-mm-dd hh:mm", 5: "0.00", 6: "0"})
    ws5 = wb.create_sheet("Notes")
    for line in ["Odor events, May 1 to September 30, 2026, under CP-14 rev. 4: a household's complaints noticed within "
                 "two hours of its first complaint are one event; the event hour is the clock hour of the time noticed, "
                 "shifted to Pacific standard time to match the station.",
                 "Complaints by month count each report once (CP-14 section 3: a report that reaches the District "
                 "more than once is one complaint), by the time the odor was noticed. App reports whose noticed time "
                 "falls after the time they were received (all before April 30) carry the time in UTC and are "
                 "converted to local time; the %d complaints carried over from the retired system on January 6 were "
                 "noticed in November and December 2025. Complaints from outside the District's boundary are left "
                 "out." % F["carried"],
                 "Station hours: an hour has a usable wind when nine of its twelve 5-minute records are wind "
                 "measurements; records taken while the vane was held for the monthly audits or the monitor was "
                 "unplugged (June 9 to 11) are not. Months are local calendar months.",
                 "Prepared by Alyssa White, October 12, 2026."]:
        ws5.append([line])
    ws5.column_dimensions["A"].width = 130
    for row in ws5.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    PP.normalize_xlsx(path, "Alyssa White", "2026-10-07T21:30:00Z", "2026-10-12T23:40:00Z")


def write(golden, F):
    os.makedirs(golden, exist_ok=True)
    workbook(os.path.join(golden, XLSX), F)
    determination(os.path.join(golden, DOCX), F)
    open(os.path.join(golden, PNG), "wb").write(chart_png(F))
