"""The shipped records, as each system exports them: the council's KS2 to KS4 transfer file and its establishment
register, the Department's performance tables and national estimates, the trust warehouse's attendance extract and
movements log, the two summer series' examination results, and the two context workbooks."""
import csv
import datetime as dt
import numpy as np
import common as K
import trust as T
import world as W

D = dt.date
YEAR_LABEL = {2022: "2021/22", 2023: "2022/23", 2024: "2023/24"}
P8_SD = {2022: 1.29, 2023: 1.27, 2024: 1.25}      # the national pupil-level spread the confidence intervals use
OUT_LA = ["350", "354", "381", "382", "383", "384", "888", "891", "925", "373", "806", "330"]


def write_csv(path, head, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        wr = csv.writer(f, lineterminator="\r\n")
        wr.writerow(head)
        for r in rows:
            wr.writerow(["" if v is None else v for v in r])


def la_estab(n):
    return int("%s%04d" % (K.LA_CODE, n))


# ------------------------------------------------------------------ identities
def primary_ids():
    """Each primary's URN and DfE number, and for the five that converted, the predecessor's."""
    out = {}
    for p, name in enumerate(K.PRIMARY_NAMES):
        faith = p in K.CATHOLIC or p in K.CHURCH
        urn = 105602 + 13 * p + (p % 5)
        estab = (3300 if faith else 2000) + p + 1
        rec = dict(urn=urn, estab=estab, name=name, faith=faith, conv=None)
        if p in K.CONVERSIONS:
            rec["conv"] = dict(date=K.CONVERSIONS[p], urn=144310 + 37 * p, estab=2100 + p,
                               name=name.replace(" Primary School", " Primary Academy"))
        out[p] = rec
    return out


def ks2_urn(pid, p, ks2_year):
    rec = pid[p]
    c = rec["conv"]
    if c and c["date"] <= D(ks2_year, 5, 13):           # the tests sat in the week of 13 May
        return c["urn"]
    return rec["urn"]


def out_of_borough(w):
    """The key stage 2 schools outside the borough: a pool of URNs and their LA numbers."""
    rg = w.r("out-of-borough schools")
    pool = []
    for i in range(70):
        pool.append((str(OUT_LA[int(rg.integers(len(OUT_LA)))]), int(100000 + rg.integers(1000, 49999))))
    return pool


# ------------------------------------------------------------------ the transfer file
def transfer(w):
    pid = primary_ids()
    pool = out_of_borough(w)
    Dd = w.D.copy()
    rg = w.r("pupil references")
    order = rg.permutation(len(Dd))
    refs = ["BTP%06d" % (310000 + i * 7 + int(rg.integers(0, 7))) for i in range(len(Dd))]
    rows = []
    for j, r in enumerate(Dd.itertuples(index=False)):
        ref = refs[order[j]]
        ks2y = K.KS2_YEAR[r.year]
        if r.no_ks2:
            la, urn, rd, ma, pa = None, None, None, None, None
        elif r.p >= 0:
            la, urn = K.LA_CODE, ks2_urn(pid, r.p, ks2y)
            rd, ma, pa = int(r.ks2_read), int(r.ks2_mat), "%.1f" % r.ks2
        else:
            rgi = W.rng_for("incomer school", K.SEED, r.year, j)
            la, urn = pool[int(rgi.integers(len(pool)))]
            rd, ma, pa = int(r.ks2_read), int(r.ks2_mat), "%.1f" % r.ks2
        unauth = int(r.unauth)
        auth = int(r.missed) - unauth
        rows.append([ref, YEAR_LABEL[r.year], K.SEC_URN[r.s], la_estab(K.SEC_ESTAB[r.s]), w.sec_name[r.s],
                     "F" if r.girl else "M", r.sen, "Y" if r.eal else "N", "Y" if r.dis else "N",
                     ks2y if not r.no_ks2 else None, la, urn, rd, ma, pa, int(r.possible), auth, unauth,
                     "%.1f" % r.a8 if r.a8 % 1 else "%d" % r.a8, 2 * r.eng, 2 * r.mat,
                     "%g" % r.ebac, "%g" % r.opn, None if r.no_ks2 else "%.2f" % r.p8])
    rows.sort(key=lambda x: (x[1], x[2], x[0]))
    head = ["pupil_ref", "ks4_year", "ks4_urn", "ks4_laestab", "ks4_school", "gender", "sen_provision", "eal",
            "disadvantaged", "ks2_year", "ks2_la", "ks2_urn", "ks2_read_scaled", "ks2_mat_scaled",
            "ks2_prior_attainment", "y11_sessions_possible", "y11_sessions_auth_absence",
            "y11_sessions_unauth_absence", "att8", "att8_english", "att8_maths", "att8_ebacc", "att8_open", "p8_score"]
    return head, rows


def register(w):
    pid = primary_ids()
    rows = []
    gov = {s[0]: s[3] for s in K.SECONDARIES}
    opened = {"HAL": "2014-04-01", "BPK": "2011-09-01", "FRN": "2015-01-01", "CDW": "2012-04-01", "MSG": "2016-06-01",
              "SLM": "2013-02-01", "RDH": "2011-08-01", "WRF": "2016-11-01", "DNH": "2012-08-01", "OKV": "2014-03-01",
              "HSK": "2017-04-01"}
    for k in w.sec:
        acad = gov[k].startswith("Academy")
        rel = "Roman Catholic" if k == "SWF" else ("Church of England" if k == "OKV" else "Does not apply")
        trust = K.TRUST if k in K.TRUST_KEYS else ({"DNH": "Pennine Ridge Learning Partnership",
                                                   "OKV": "Brantholme Diocesan Academy Trust",
                                                   "HSK": "Northfold Schools Trust"}.get(k))
        rows.append([K.SEC_URN[k], la_estab(K.SEC_ESTAB[k]), w.sec_name[k], "Secondary", gov[k], rel, trust, "Open",
                     opened.get(k, "1974-04-01"), None, 105800 + K.SEC_ESTAB[k] % 100 if acad else None])
    for p, rec in pid.items():
        rel = "Roman Catholic" if p in K.CATHOLIC else ("Church of England" if p in K.CHURCH else "Does not apply")
        typ = "Voluntary aided school" if p in K.CATHOLIC else ("Voluntary controlled school" if p in K.CHURCH
                                                                 else "Community school")
        c = rec["conv"]
        if c:
            close = c["date"] - dt.timedelta(days=1)
            rows.append([rec["urn"], la_estab(rec["estab"]), rec["name"], "Primary", typ, rel, None, "Closed",
                         "1974-04-01", close.isoformat(), None])
            rows.append([c["urn"], la_estab(c["estab"]), c["name"], "Primary", "Academy converter", rel,
                         "Brantholme Primary Partnership Trust" if p % 2 else "Moorside Learning Trust", "Open",
                         c["date"].isoformat(), None, rec["urn"]])
        else:
            rows.append([rec["urn"], la_estab(rec["estab"]), rec["name"], "Primary", typ, rel, None, "Open",
                         "1974-04-01", None, None])
    rows += [[106041, la_estab(7002), "Brantholme Hollies School", "Not applicable", "Community special school",
              "Does not apply", None, "Open", "1974-04-01", None, None],
             [106044, la_estab(7005), "Ridgeway Park School", "Not applicable", "Foundation special school",
              "Does not apply", None, "Open", "1981-09-01", None, None],
             [131928, la_estab(1104), "Brantholme Learning Centre", "Not applicable", "Pupil referral unit",
              "Does not apply", None, "Open", "2003-09-01", None, None]]
    rows.sort(key=lambda r: (r[3] != "Secondary", r[3] != "Primary", r[2], r[7]))
    head = ["urn", "laestab", "establishment_name", "phase", "establishment_type", "religious_character",
            "trust_name", "status", "open_date", "close_date", "predecessor_urn"]
    return head, rows


def tables(w):
    Dd = w.D
    rows = []
    for k in w.sec:
        for y in K.COHORTS:
            g = Dd[(Dd.s == k) & (Dd.year == y)]
            gp = g[~g.p8.isna()]
            m = gp.p8.mean()
            half = 1.96 * P8_SD[y] / np.sqrt(len(gp))
            rows.append([K.SEC_URN[k], w.sec_name[k], YEAR_LABEL[y], len(g), len(gp), "%.2f" % m, "%.2f" % (m - half),
                         "%.2f" % (m + half), "%.1f" % g.a8.mean(), "%d" % round(100 * g.dis.mean())])
    head = ["urn", "school_name", "year", "pupils_end_ks4", "pupils_with_p8", "p8_score", "p8_lower_ci",
            "p8_upper_ci", "att8_average", "pct_disadvantaged"]
    return head, rows


def estimates(path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    first = True
    for y in K.COHORTS:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = YEAR_LABEL[y].replace("/", "-")
        ws.append(["Prior attainment group", "KS2 average scaled score from", "to below", "Estimated Attainment 8",
                   "English element", "Maths element", "EBacc element", "Open element"])
        for r in W.pag_table(y):
            ws.append([r["pag"], r["low"], r["high"], r["a8"], r["eng"], r["mat"], r["ebac"], r["open"]])
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDE5EC")
        ws.freeze_panes = "A2"
        for col, wd in zip("ABCDEFGH", (22, 28, 11, 22, 16, 14, 15, 13)):
            ws.column_dimensions[col].width = wd
    ws = wb.create_sheet("Notes")
    for line in ["Estimated Attainment 8 by key stage 2 prior attainment group, the national averages each pupil's",
                 "Progress 8 is measured against: Progress 8 = (Attainment 8 - estimated Attainment 8) / 10.",
                 "Prior attainment is the average of the key stage 2 reading and mathematics scaled scores.",
                 "Final (revised) estimates for each year; the upper bound of each group is exclusive."]:
        ws.append([line])
    ws.column_dimensions["A"].width = 100
    wb.properties.creator = "Department for Education"
    wb.save(path)


# ------------------------------------------------------------------ the trust's records
def attendance(tr):
    pupils = {p["id"]: p for p in tr.pupils}
    rows = []
    names = {k: K.SECONDARIES[[s[0] for s in K.SECONDARIES].index(k)][1] for k in K.TRUST_KEYS}
    for a in tr.att:
        p = pupils[a["id"]]
        rows.append([K.EXTRACT_DATE.isoformat(), names[a["academy_now"]], a["adm"], a["upn"], a["yg_now"],
                     "Y" if a["status"] == "On roll" else "N", "2025/26", HT_NAME[a["half"]], a["possible"],
                     a["present"], a["aea"], a["auth"], a["unauth"]])
    order = {n: i for i, n in enumerate(HT_NAME)}
    rows.sort(key=lambda r: (r[1], r[3], order[r[7]], r[2]))
    head = ["extract_date", "academy", "admission_no", "upn", "year_group", "on_roll", "academic_year", "half_term",
            "sessions_possible", "sessions_present", "sessions_approved_activity", "sessions_auth_absence",
            "sessions_unauth_absence"]
    return head, rows


HT_NAME = ["Autumn 1", "Autumn 2", "Spring 1", "Spring 2", "Summer 1", "Summer 2"]


def movements(tr):
    names = {k: K.SECONDARIES[[s[0] for s in K.SECONDARIES].index(k)][1] for k in K.TRUST_KEYS}
    rows = []
    for m in tr.moves:
        d = T.HT_START[m["half"]]
        rows.append([d.isoformat(), m["upn"], m["kind"], names.get(m["from_k"], "Outside the trust"),
                     names.get(m["to_k"], "Outside the trust"), m["reason"]])
    rows.sort(key=lambda r: (r[0], r[2], r[1]))
    head = ["movement_date", "upn", "movement", "from_school", "to_school", "reason"]
    return head, rows


CENTRE = {"HAL": 31842, "BPK": 31207, "FRN": 31655, "CDW": 31318, "MSG": 31730, "SLM": 31491, "RDH": 31164,
          "WRF": 31903}


def results(tr, series):
    names = {k: K.SECONDARIES[[s[0] for s in K.SECONDARIES].index(k)][1] for k in K.TRUST_KEYS}
    rs = [r for r in tr.res if r["series"] == series]
    cand = {}
    for r in sorted(rs, key=lambda r: (r["academy"], r["upn"])):
        key = (r["academy"], r["upn"])
        if key not in cand:
            cand[key] = 1000 + len([1 for c in cand if c[0] == r["academy"]]) + 1
    rows = []
    for r in rs:
        title, code, kind = T.SUBJECTS[r["subj"]]
        qual = "Technical Award" if kind == "TECH" else "GCSE (9-1)"
        rows.append([series, CENTRE[r["academy"]], names[r["academy"]], cand[(r["academy"], r["upn"])], r["upn"],
                     r["yg"], qual, title, code, r["grade"]])
    rows.sort(key=lambda x: (x[1], x[3], x[7]))
    head = ["series", "centre_number", "centre_name", "candidate_number", "upn", "year_group", "qualification",
            "subject", "discount_code", "grade"]
    return head, rows


# ------------------------------------------------------------------ the context workbooks (distractors)
def workforce(path, w):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    rg = w.r("workforce")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Secondary schools"
    ws.append(["URN", "School", "Teachers (FTE)", "Teachers (headcount)", "Pupil to teacher ratio",
               "Teaching assistants (FTE)", "Teacher leavers 2024/25 (% of FTE)", "Vacancies (Nov 2025)",
               "Teachers with QTS (%)", "Average teacher absence (days)"])
    for k in w.sec:
        roll = K.PLANNED[k] * 5
        fte = roll / rg.uniform(15.2, 17.4)
        ws.append([K.SEC_URN[k], w.sec_name[k], round(fte, 1), int(fte * rg.uniform(1.05, 1.13)),
                   round(roll / fte, 1), round(fte * rg.uniform(0.28, 0.42), 1), round(rg.uniform(8.5, 15.5), 1),
                   int(rg.integers(0, 5)), round(rg.uniform(93.0, 99.5), 1), round(rg.uniform(3.6, 6.9), 1)])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E3E9EF")
    ws.freeze_panes = "C2"
    for col, wd in zip("ABCDEFGHIJ", (9, 38, 14, 18, 20, 22, 30, 18, 20, 26)):
        ws.column_dimensions[col].width = wd
    wb.properties.creator = K.COUNCIL
    wb.save(path)


def funding(path, w):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    rg = w.r("funding")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2025-26 allocations"
    ws.append(["URN", "School", "Pupils on roll (Oct 2024 census)", "Core schools budget (GBP)",
               "Funding per pupil (GBP)", "Pupil premium (GBP)", "High needs top-up (GBP)", "Notes"])
    for k in w.sec:
        roll = int(K.PLANNED[k] * 5 * rg.uniform(0.94, 1.03))
        per = rg.uniform(6420, 7310)
        pp = int(roll * rg.uniform(0.18, 0.46) * 1075)
        ws.append([K.SEC_URN[k], w.sec_name[k], roll, int(roll * per), int(round(per)), pp,
                   int(rg.uniform(40, 260) * 1000), "Academy: allocation via ESFA" if k not in ("SWF", "TRS", "KSB")
                   else "Maintained: allocation via the council"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E3E9EF")
    ws.freeze_panes = "C2"
    for col, wd in zip("ABCDEFGH", (9, 38, 28, 24, 22, 18, 22, 36)):
        ws.column_dimensions[col].width = wd
    wb.properties.creator = K.COUNCIL
    wb.save(path)
