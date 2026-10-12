"""Write the data files of target/: the three sittings, the register, the census, the service directory, the
2027 visit plan, the road distances and the two neighbouring files (specialist rates, PL register)."""
import csv
import datetime as dt
import os

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import core
import postprocess as PP
import world as WD

ORG = "Catholic Schools Office"
FILES = {
    "r2024": "naplan2024_legacy_extract.csv",
    "r2025": "NAPLAN2025_results_portal.xlsx",
    "r2026": "naplan_results_2026.csv",
    "register": "school_register_2026.xlsx",
    "census": "census_aug2026_enrolments.csv",
    "directory": "facilities_service_directory_feb2026.xlsx",
    "plan": "hub_visit_plan_2027.xlsx",
    "road": "hub_school_road_distances.csv",
    "standard": "LT-AS-04_analysing_school_results.pdf",
    "guidelines": "reading_partnership_2027_guidelines.docx",
    "notes": "data_team_notes.txt",
    "rates": "literacy_specialist_rates_2027.xlsx",
    "pl": "literacy_PL_register_2026.csv",
}
DISTRACTORS = ["rates", "pl"]
TEST_NAME = {"RD": "Reading", "WR": "Writing", "SP": "Spelling", "GP": "Grammar and Punctuation", "NU": "Numeracy"}
PART_NAME = {"P": "participated", "A": "absent", "E": "exempt", "W": "withdrawn", "S": "sanctioned abandonment"}


def school_names(W):
    rng = np.random.default_rng([core.SITTINGS[0], 7])
    out = []
    for t in W["title"]:
        kind = "Catholic Primary School" if rng.random() < 0.18 else "Primary School"
        out.append("%s %s" % (t, kind))
    return out


def hub_names(W):
    return {h: W["hubs"][h]["locality"] for h in "ABCDE"}


def _write_csv(path, header, rows):
    with open(path, "w", newline="") as f:
        wr = csv.writer(f, lineterminator="\r\n")
        wr.writerow(header)
        wr.writerows(rows)


def _fmt(x, dec):
    if x != x:
        return ""
    return ("%.1f" % x) if dec else ("%d" % int(round(x)))


# ------------------------------------------------------------------------------------------- results
def r2024(path, W, df):
    codes = W["codes"]
    rows = []
    for s, st, yl, t, pc, sc in zip(df.school.to_numpy(), df.student.to_numpy(), df.year.to_numpy(),
                                    df.test.to_numpy(), df.part.to_numpy(), df.score.to_numpy()):
        rows.append(["20240903", codes[s], "%d" % st, "%02d" % yl, t, pc, _fmt(sc, True)])
    order = sorted(range(len(rows)), key=lambda i: (rows[i][1], rows[i][3], rows[i][2], "RDWRSPGPNU".index(rows[i][4])))
    _write_csv(path, ["EXTRACT_DT", "SCHL_CD", "STU_ID", "YR_LVL", "TEST_CD", "PART_CD", "SCALE_SCR"],
               [rows[i] for i in order])


def r2025(path, W, df, names):
    wb = Workbook(write_only=True)
    codes = W["codes"]
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    for yl in (3, 5):
        ws = wb.create_sheet("Year %d" % yl)
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 40
        ws.column_dimensions["C"].width = 16
        ws.column_dimensions["D"].width = 26
        ws.column_dimensions["E"].width = 14
        ws.column_dimensions["F"].width = 12
        ws.append(["NAPLAN 2025 student results, Year %d" % yl])
        ws.append(["Downloaded from the results portal 22/08/2025 by the Catholic Schools Office data team"])
        ws.append(["Participation: P participated, A absent, E exempt, W withdrawn, S sanctioned abandonment"])
        ws.append([])
        ws.append(["School Code", "School", "Student Number", "Domain", "Participation", "Scale Score"])
        d = df[df.year == yl]
        idx = sorted(range(len(d)), key=lambda i: (codes[d.school.iat[i]], d.student.iat[i],
                                                    "RDWRSPGPNU".index(d.test.iat[i])))
        sch, stu, tst, prt, sco = (d.school.to_numpy(), d.student.to_numpy(), d.test.to_numpy(), d.part.to_numpy(),
                                   d.score.to_numpy())
        for i in idx:
            v = sco[i]
            ws.append([codes[sch[i]], names[sch[i]], int(stu[i]), TEST_NAME[tst[i]], prt[i],
                       None if v != v else int(round(v))])
    wb.save(path)
    PP.normalize_xlsx(path, "Joshua Nguyen", "2025-08-22T01:14:00Z", "2025-08-22T01:31:00Z", company=ORG)


def r2026(path, W, df):
    codes = W["codes"]
    rows = []
    for s, st, yl, t, pc, sc in zip(df.school.to_numpy(), df.student.to_numpy(), df.year.to_numpy(),
                                    df.test.to_numpy(), df.part.to_numpy(), df.score.to_numpy()):
        rows.append([codes[s], "%d" % st, "Year %d" % yl, TEST_NAME[t].lower(), PART_NAME[pc], _fmt(sc, False),
                     "2026-03"])
    order = sorted(range(len(rows)), key=lambda i: (rows[i][0], rows[i][2], rows[i][1], rows[i][3]))
    _write_csv(path, ["school_id", "student_ref", "year_level", "domain", "participation_status", "scaled_score",
                      "test_window"], [rows[i] for i in order])


# ------------------------------------------------------------------------------------------- workbooks
def _sheet(wb, title, header, rows, widths, note=None, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    r0 = 1
    if note:
        for line in note:
            ws.cell(row=r0, column=1, value=line)
            r0 += 1
        r0 += 1
    for j, h in enumerate(header, start=1):
        c = ws.cell(row=r0, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
    for i, row in enumerate(rows, start=r0 + 1):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=v)
            if isinstance(v, dt.date):
                c.number_format = "dd/mm/yyyy"
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    return ws


def register(path, W, names):
    rows = []
    reg = W["region"]
    rng = np.random.default_rng([2026, 3])
    for i in range(len(names)):
        rows.append([W["codes"][i], names[i], W["loc"][i], WD.REGION_NAME[reg[i]], "Primary", "F-6", "Open",
                     int(rng.integers(1885, 1999))])
    for k, s in enumerate(W["sec"]):
        rows.append([W["sec_codes"][k], s["name"], s["locality"], WD.REGION_NAME[s["region"]], "Secondary", "7-12",
                     "Open", int(rng.integers(1902, 2012))])
    rows.sort(key=lambda r: r[0])
    wb = Workbook()
    _sheet(wb, "Schools", ["School code", "School name", "Locality", "Region", "School type", "Year levels",
                           "Status", "Established"], rows, [12, 42, 22, 20, 12, 11, 8, 12], first=True,
           note=["Catholic Schools Office school register", "Refreshed 3 March 2026 (system data team)"])
    wb.save(path)
    PP.normalize_xlsx(path, "Joshua Nguyen", "2019-02-11T03:20:00Z", "2026-03-03T00:47:00Z", company=ORG)


def directory(path, W, names, hubs):
    rows = []
    for i in range(len(names)):
        rows.append([W["codes"][i], names[i], "Primary school", W["loc"][i], hubs[W["dir_hub"][i]]])
    road = W["road"]
    for k, s in enumerate(W["sec"]):
        d = [np.hypot(*(s["xy"] - W["hubs"][h]["xy"])) for h in "ABCD"]
        rows.append([W["sec_codes"][k], s["name"], "Secondary college", s["locality"], hubs["ABCD"[int(np.argmin(d))]]])
    rows.sort(key=lambda r: r[0])
    wb = Workbook()
    _sheet(wb, "Directory", ["Site code", "Site name", "Site type", "Locality", "Servicing hub"], rows,
           [10, 42, 18, 22, 20], first=True,
           note=["Facilities service directory, all sites", "Annual refresh 16 February 2026, Facilities (Alice Wood)"])
    hs = [[hubs[h], W["hubs"][h]["locality"], "Opened July 2026" if h == "E" else ""] for h in "ABCDE"]
    _sheet(wb, "Hubs", ["Hub", "Base locality", "Remarks"], hs, [20, 22, 20])
    wb.save(path)
    PP.normalize_xlsx(path, "Alice Wood", "2021-02-08T22:41:00Z", "2026-08-03T05:12:00Z", company=ORG)


ABBR = [("North ", "Nth "), ("South ", "Sth "), (" Heights", " Hts"), ("Upper ", "Upr ")]
# each hub's coordinator keeps the list in their own way: (share written with " PS", " Primary",
# " Primary School"; share with the apostrophe dropped; share with "St" written out; share of localities
# abbreviated). One or two entries per sheet break the coordinator's habit, as hand-kept lists do.
STYLE = {"A": ((0.88, 0.0, 0.0), 0.0, 0.0, 0.92), "B": ((0.0, 0.06, 0.0), 0.04, 0.0, 0.0),
         "C": ((0.0, 0.83, 0.0), 0.9, 0.0, 0.3), "D": ((0.05, 0.0, 0.7), 0.0, 0.35, 0.6),
         "E": ((0.45, 0.0, 0.0), 0.5, 0.0, 0.0)}


def plan_label(title, loc, rng, style):
    suf, drop, saint, abbr = style
    r = rng.random()
    nm = title + (" PS" if r < suf[0] else " Primary" if r < suf[0] + suf[1] else
                  " Primary School" if r < sum(suf) else "")
    if rng.random() < drop:
        nm = nm.replace("'", "")
    if nm.startswith("St ") and rng.random() < saint:
        nm = "Saint " + nm[3:]
    if rng.random() < abbr:
        for a, b in ABBR:
            loc = loc.replace(a, b)
    return nm, loc


def plan(path, W, hubs):
    rng = np.random.default_rng([2027, 1])
    wb = Workbook()
    first = True
    for h in "ABCDE":
        idx = [i for i in range(len(W["plan"])) if W["plan"][i] == h]
        idx.sort(key=lambda i: (W["loc"][i], W["title"][i]))
        rows = []
        for i in idx:
            nm, lc = plan_label(W["title"][i], W["loc"][i], rng, STYLE[h])
            rows.append([nm, lc, ["Week A", "Week B"][int(rng.integers(2))],
                         ["Mon", "Tue", "Wed", "Thu", "Fri"][int(rng.integers(5))]])
        _sheet(wb, hubs[h], ["School", "Locality", "Fortnight", "Visit day"], rows, [34, 22, 11, 10], first=first,
               note=["2027 school servicing plan, %s hub" % hubs[h],
                     "Hub coordinators, version of 6 November 2026; takes effect from the start of Term 1 2027"])
        first = False
    wb.save(path)
    PP.normalize_xlsx(path, "Diana Hardy", "2026-10-19T23:05:00Z", "2026-11-06T04:38:00Z", company=ORG)


def census(path, W):
    rows = []
    c = W["census"]
    rolled = W["rolled"]
    for i in range(len(W["codes"])):
        for li, lv in enumerate(W["levels"]):
            # a school already moved to the 2027 school year in the student system: every student one level up
            held = str(li + 1) if rolled[i] else lv
            rows.append(["07/08/2026", int(W["codes"][i]), held, int(c[i, li])])
    rng = np.random.default_rng([2026, 8])
    for k, s in enumerate(W["sec"]):
        base = rng.uniform(80, 210)
        for lv in range(7, 13):
            rows.append(["07/08/2026", int(W["sec_codes"][k]), str(lv),
                         int(max(10, rng.poisson(base * (0.97 if lv > 10 else 1.0))))])
    rows.sort(key=lambda r: (r[1], ["F", "1", "2", "3", "4", "5", "6", "7", "8", "9", "10", "11", "12"].index(r[2])))
    _write_csv(path, ["census_date", "school_no", "year_level", "headcount"], rows)


def road(path, W, hubs):
    rows = []
    rng = np.random.default_rng([2026, 9])
    metro = np.array([r in core.METRO for r in core.np.array(W["w"]["region"])])
    for k, h in enumerate("ABCDE"):
        for i in range(len(W["codes"])):
            km = int(W["road"][i, k])
            speed = 38.0 if (metro[i] and km < 60) else 78.0
            mins = int(round(km / speed * 60 + rng.uniform(2, 7)))
            if metro[i] and h not in "AB":
                continue          # metropolitan schools are routed to the two metropolitan hubs only
            rows.append([hubs[h], int(W["codes"][i]), km, mins])
    rows.sort(key=lambda r: (r[0], r[1]))
    _write_csv(path, ["hub", "school_no", "road_km", "drive_min"], rows)


def rates(path):
    wb = Workbook()
    rows = [["Specialist reading teacher", "Level 1", 742.00, 0.96], ["Specialist reading teacher", "Level 2", 798.00, 0.96],
            ["Lead specialist", "Level 3", 865.00, "Fleet vehicle"], ["Relief specialist (casual)", "Casual", 688.50, 0.88]]
    _sheet(wb, "Day rates", ["Role", "Classification", "Day rate (AUD, ex GST)", "Vehicle allowance (AUD per km)"],
           rows, [30, 14, 22, 28], first=True,
           note=["Specialist literacy services: schedule of rates, 1 January to 31 December 2027",
                 "Provided by the panel provider with its tender response of 30 September 2026"])
    _sheet(wb, "Conditions", ["Item", "Condition"], [
        ["Minimum engagement", "Half a day per school visit"],
        ["Travel time", "Paid at the day rate pro rata beyond 90 minutes' travel each way"],
        ["Overnight stays", "Accommodation and meals reimbursed at the Catholic Schools Office travel rates"],
        ["Indexation", "Rates are fixed for the 2027 school year"],
    ], [24, 80])
    wb.save(path)
    PP.normalize_xlsx(path, "Laura Proctor", "2026-09-30T06:02:00Z", "2026-10-02T00:16:00Z", company=ORG)


def pl_register(path, W):
    rng = np.random.default_rng([2026, 12])
    courses = [("LIT-104", "Explicit teaching of phonics, F-2", 12), ("LIT-118", "Reading fluency in the middle years", 9),
               ("LIT-121", "Comprehension strategies, Years 3-6", 9), ("LIT-130", "Writing conferences", 6),
               ("LIT-142", "Using NAPLAN reports in planning", 3), ("LIT-150", "Morphology and spelling", 6),
               ("LIT-163", "Oral language in the early years", 6)]
    rows = []
    S = len(W["codes"])
    staff = rng.choice(np.arange(200000, 899999), size=6000, replace=False)
    k = 0
    for i in range(S):
        m = int(rng.poisson(9))
        for _ in range(m):
            c = courses[int(rng.integers(len(courses)))]
            start = dt.date(2026, 1, 28) + dt.timedelta(days=int(rng.integers(0, 260)))
            status = "Completed" if start < dt.date(2026, 9, 1) else ["Completed", "Enrolled", "Withdrawn"][int(rng.integers(3))]
            rows.append([W["codes"][i], int(staff[k % len(staff)]), c[0], c[1], start.strftime("%d/%m/%Y"), c[2], status])
            k += 1
    rows.sort(key=lambda r: (r[0], r[4][6:] + r[4][3:5] + r[4][:2], r[1]))
    _write_csv(path, ["school_code", "staff_no", "course_code", "course_title", "start_date", "hours", "status"], rows)


def write(W, tables, target):
    os.makedirs(target, exist_ok=True)
    names = school_names(W)
    hubs = hub_names(W)
    P = lambda k: os.path.join(target, FILES[k])
    r2024(P("r2024"), W, tables[0])
    r2025(P("r2025"), W, tables[1], names)
    r2026(P("r2026"), W, tables[2])
    register(P("register"), W, names)
    census(P("census"), W)
    directory(P("directory"), W, names, hubs)
    plan(P("plan"), W, hubs)
    road(P("road"), W, hubs)
    rates(P("rates"))
    pl_register(P("pl"), W)
    return names, hubs
