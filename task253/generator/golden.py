"""The two golden deliverables: reading_partnership_2027.xlsx and reading_partnership_2027_hubs.png."""
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import asks as A
import core
import postprocess as PP
import world as WD

XLSX = "reading_partnership_2027.xlsx"
PNG = "reading_partnership_2027_hubs.png"
NAVY = "1F3864"


def _head(ws, row, labels, widths=None):
    thin = Side(style="thin", color="8EA9DB")
    for j, h in enumerate(labels, start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center" if j > 1 else "left", vertical="center", wrap_text=True)
        c.border = Border(bottom=thin)
    if widths:
        for j, w in enumerate(widths):
            ws.column_dimensions[chr(65 + j)].width = w


def _total(ws, row, ncol):
    for j in range(1, ncol + 1):
        c = ws.cell(row=row, column=j)
        c.font = Font(bold=True)
        c.border = Border(top=Side(style="thin", color="000000"))


def workbook(path, W, fd, figs, names, hubs):
    p = fd["p"]
    reg = W["region"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Places"
    ws["A1"] = "Reading Partnership 2027: places to commission"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "For the executive meeting, 18 November 2026. NAPLAN Year 5 reading, 2024 to 2026 sittings, LT-AS-04 model."
    ws["A2"].font = Font(italic=True, color="595959")
    ws["A4"] = "Places to commission"
    ws["A4"].font = Font(bold=True)
    ws["B4"] = figs["places"]
    ws["B4"].font = Font(bold=True, size=14)
    ws["A5"] = "Sum over the 264 primary schools of each school's chance that its underlying Year 5 reading average is below 480"
    ws["A5"].font = Font(color="595959", size=9)
    _head(ws, 7, ["Region", "Primary schools", "Share of places (%)"], [44, 16, 20])
    r = 8
    for k in core.REGIONS:
        ws.cell(row=r, column=1, value=WD.REGION_NAME[k])
        ws.cell(row=r, column=2, value=int((reg == k).sum()))
        c = ws.cell(row=r, column=3, value=round(figs["shares"][k], 1))
        c.number_format = "0.0"
        r += 1
    ws.cell(row=r, column=1, value="All regions")
    ws.cell(row=r, column=2, value=len(reg))
    c = ws.cell(row=r, column=3, value=100.0)
    c.number_format = "0.0"
    _total(ws, r, 3)
    ws.freeze_panes = "A8"

    ws = wb.create_sheet("Regions")
    ws["A1"] = "Year 3 to 6 students the places will serve in 2027, by region"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("2027 Years 3 to 6: the August 2026 census's Years 2 to 5, or its Years 3 to 6 at the %d schools "
                "already rolled over to 2027 in the student system." % int(W["rolled"].sum()))
    ws["A2"].font = Font(italic=True, color="595959")
    _head(ws, 4, ["Region", "Students served"], [44, 18])
    r = 5
    for k in core.REGIONS:
        ws.cell(row=r, column=1, value=WD.REGION_NAME[k])
        ws.cell(row=r, column=2, value=int(round(figs["students_region"][k]))).number_format = "#,##0"
        r += 1
    ws.cell(row=r, column=1, value="All regions")
    ws.cell(row=r, column=2, value=int(round(figs["students_total"]))).number_format = "#,##0"
    _total(ws, r, 2)

    ws = wb.create_sheet("Hubs")
    ws["A1"] = "Hubs: students served in 2027 and one fortnightly round of visits"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Hubs as in the coordinators' 2027 servicing plan; road km hub to school and back."
    ws["A2"].font = Font(italic=True, color="595959")
    _head(ws, 4, ["Hub", "Students served", "Fortnightly round (km)"], [26, 18, 24])
    r = 5
    for h in "ABCDE":
        ws.cell(row=r, column=1, value=hubs[h])
        ws.cell(row=r, column=2, value=int(round(figs["students_hub"][h]))).number_format = "#,##0"
        ws.cell(row=r, column=3, value=int(round(figs["km_hub"][h]))).number_format = "#,##0"
        r += 1
    ws.cell(row=r, column=1, value="All hubs")
    ws.cell(row=r, column=2, value=int(round(figs["students_total"]))).number_format = "#,##0"
    ws.cell(row=r, column=3, value=int(round(figs["km_total"]))).number_format = "#,##0"
    _total(ws, r, 3)

    ws = wb.create_sheet("Schools")
    _head(ws, 1, ["School code", "School", "Locality", "Region", "2027 hub", "Year 5 readers 2024-26",
                  "Model average", "Chance below 480", "Years 3-6 in 2027", "One-way km"],
          [11, 40, 22, 20, 18, 12, 11, 11, 11, 10])
    stu = A.students(W["census"])
    hk = {h: i for i, h in enumerate("ABCDE")}
    order = sorted(range(len(p)), key=lambda i: W["codes"][i])
    for r, i in enumerate(order, start=2):
        vals = [W["codes"][i], names[i], W["loc"][i], WD.REGION_NAME[reg[i]], hubs[W["plan"][i]],
                int(fd["n"][i].sum()), round(float(fd["mean"][i]), 1), round(float(p[i]), 3), int(stu[i]),
                int(W["road"][i, hk[W["plan"][i]]])]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            if j == 7:
                c.number_format = "0.0"
            if j == 8:
                c.number_format = "0.000"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = "A1:J%d" % (len(order) + 1)
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws = wb.create_sheet("Notes")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 96
    notes = [
        ("As at", "10 November 2026"),
        ("Results", "NAPLAN Year 5 reading, participating students only (LT-AS-04 2.1): 2024 legacy extract, "
                    "2025 results-portal download, 2026 data-platform export."),
        ("Model", "LT-AS-04 3.1 to 3.4: school and year-group random effects, student residual, REML, "
                  "calculations at the fitted values."),
        ("Places", "Guidelines 2 and 3. Each school's chance of an underlying average below 480 is added over "
                   "the 264 primary schools; the total is rounded once to whole places."),
        ("Shares", "Each region's part of that total, per cent to one decimal."),
        ("Students", "2027 Years 3 to 6, weighted by each school's chance: the August 2026 census's Years 2 to 5, "
                     "everyone moving up a level for 2027. %d primary schools show Year 1 to Year 7 and no "
                     "Foundation in the census run of 4 November: their student system had already been rolled "
                     "over to 2027, so their Years 3 to 6 are taken as they stand. Census school numbers padded "
                     "to four characters to match the register." % int(W["rolled"].sum())),
        ("Hubs", "From the hub coordinators' 2027 servicing plan (6 November 2026), matched on school name and "
                 "locality; the February facilities directory predates the 2027 hub changes."),
        ("Kilometres", "Guidelines 4: hub to school and back for each visit; routing service one-way km doubled, "
                       "weighted by each school's chance."),
        ("Prepared by", "Laura Proctor, Partnership lead, with the system data team"),
    ]
    for r, (a, b) in enumerate(notes, start=1):
        ws.cell(row=r, column=1, value=a).font = Font(bold=True)
        c = ws.cell(row=r, column=2, value=b)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    PP.normalize_xlsx(path, "Laura Proctor", "2026-11-10T00:12:00Z", "2026-11-12T05:47:00Z",
                      company="Catholic Schools Office")


def chart(path, figs, hubs):
    km = figs["km_hub"]
    st = figs["students_hub"]
    order = figs["chart_order"]
    lead = order[0]
    fig, ax = plt.subplots(figsize=(8.6, 4.4), dpi=150)
    y = np.arange(len(order))[::-1]
    vals = [round(km[h]) for h in order]
    cols = ["#1F3864" if h == lead else "#8EA9DB" for h in order]
    ax.barh(y, vals, color=cols, height=0.62)
    for yy, h, v in zip(y, order, vals):
        ax.text(v + 25, yy, "{:,} km  |  {:,} students".format(v, int(round(st[h]))), va="center", fontsize=9,
                color="#262626")
    ax.set_yticks(y)
    ax.set_yticklabels(["%s hub" % hubs[h] for h in order], fontsize=9.5)
    ax.set_xlim(0, max(vals) * 1.42)
    ax.set_xlabel("Road km, one fortnightly round of visits (hub to each partner school and back)", fontsize=9)
    ax.spines[["top", "right"]].set_visible(False)
    ax.tick_params(axis="x", labelsize=8.5)
    ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda x, _: "{:,.0f}".format(x)))
    ax.set_title("%s hub has the longest fortnightly round: %s km" % (hubs[lead], "{:,}".format(round(km[lead]))),
                 loc="left", fontsize=11.5, fontweight="bold")
    fig.text(0.01, 0.01, "Reading Partnership 2027. Students: Years 3 to 6 in 2027 at the hub's partner schools. "
             "Source: CSO data team, 2026 census and 2027 servicing plan.", fontsize=7, color="#595959")
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)


def write(golden_dir, W, fd, figs, names, hubs):
    os.makedirs(golden_dir, exist_ok=True)
    workbook(os.path.join(golden_dir, XLSX), W, fd, figs, names, hubs)
    chart(os.path.join(golden_dir, PNG), figs, hubs)
