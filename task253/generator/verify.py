"""Independent verifier: reads target/ only and recomputes the call and every graded figure on its own code
path (its own parsers, REML by EM on the restricted likelihood, its own joins).

    python3 task253/generator/verify.py [target_dir]
"""
import csv
import math
import os
import re
import sys

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
TARGET = os.path.join(os.path.dirname(HERE), "target")


def _phi(z):
    return 0.5 * (1.0 + math.erf(z / math.sqrt(2.0)))


def year5_reading(t):
    rows = []                                               # (code, sitting, score)
    with open(os.path.join(t, "naplan2024_legacy_extract.csv"), newline="") as f:
        for r in csv.DictReader(f):
            if r["YR_LVL"] == "05" and r["TEST_CD"] == "RD" and r["PART_CD"] == "P":
                rows.append((r["SCHL_CD"], 2024, float(r["SCALE_SCR"])))
    wb = load_workbook(os.path.join(t, "NAPLAN2025_results_portal.xlsx"), read_only=True)
    ws = wb["Year 5"]
    hdr = None
    for row in ws.iter_rows(values_only=True):
        if hdr is None:
            if row and row[0] == "School Code":
                hdr = {h: k for k, h in enumerate(row)}
            continue
        if row[hdr["Domain"]] == "Reading" and row[hdr["Participation"]] == "P":
            rows.append((str(row[hdr["School Code"]]), 2025, float(row[hdr["Scale Score"]])))
    with open(os.path.join(t, "naplan_results_2026.csv"), newline="") as f:
        for r in csv.DictReader(f):
            if r["year_level"] == "Year 5" and r["domain"] == "reading" and r["participation_status"] == "participated":
                rows.append((r["school_id"], 2026, float(r["scaled_score"])))
    return rows


def register(t):
    wb = load_workbook(os.path.join(t, "school_register_2026.xlsx"), read_only=True)
    out, hdr = {}, None
    for row in wb["Schools"].iter_rows(values_only=True):
        if hdr is None:
            if row and row[0] == "School code":
                hdr = {h: k for k, h in enumerate(row)}
            continue
        if row[hdr["School type"]] == "Primary":
            out[str(row[hdr["School code"]])] = {"name": row[hdr["School name"]], "loc": row[hdr["Locality"]],
                                                  "region": row[hdr["Region"]], "levels": row[hdr["Year levels"]]}
    return out


def reml_em(groups, SSW, N, tol=1e-13, maxit=400000):
    """groups: per school, a list of (n, mean). EM on the restricted likelihood (flat prior on the mean)."""
    G = sum(len(g) for g in groups)
    S = len(groups)
    tu, tv, s2 = 300.0, 100.0, SSW / (N - G)
    for it in range(maxit):
        # GLS mean and its variance
        num = den = 0.0
        cache = []
        for g in groups:
            w = [1.0 / (tv + s2 / n) for n, _ in g]
            W = sum(w)
            a = W / (1.0 + tu * W)
            b = sum(wi * y for wi, (_, y) in zip(w, g))
            num += b - tu * W * b / (1.0 + tu * W)
            den += a
            cache.append((w, W))
        mu = num / den
        vmu = 1.0 / den
        Eu2 = Ev2 = Ee = 0.0
        for g, (w, W) in zip(groups, cache):
            P = 1.0 / tu + W
            m = sum(wi * (y - mu) for wi, (_, y) in zip(w, g)) / P
            dm = -W / P
            Eu2 += 1.0 / P + m * m + dm * dm * vmu
            for wi, (n, y) in zip(w, g):
                k = (n / s2) / (1.0 / tv + n / s2)
                cv = 1.0 / (1.0 / tv + n / s2)
                r = y - mu
                ev = k * (r - m)
                dev = k * (-1.0 - dm)
                Ev2 += cv + k * k / P + ev * ev + dev * dev * vmu
                ee = (1 - k) * (r - m)
                dee = (1 - k) * (-1.0 - dm)
                Ee += n * ((1 - k) ** 2 / P + cv + ee * ee + dee * dee * vmu)
        ntu, ntv, ns2 = Eu2 / S, Ev2 / G, (SSW + Ee) / N
        done = abs(ntu - tu) / tu < tol and abs(ntv - tv) / tv < tol and abs(ns2 - s2) / s2 < tol
        tu, tv, s2 = ntu, ntv, ns2
        if done:
            break
    return tu, tv, s2, mu, it


def _norm_name(s):
    s = s.lower().replace("’", "'").replace("'", "")
    s = re.sub(r"\b(catholic primary school|primary school|primary|ps)\b", "", s)
    s = re.sub(r"\bsaint\b", "st", s)
    return " ".join(s.split())


def _norm_loc(s):
    s = " " + s.lower() + " "
    for a, b in ((" nth ", " north "), (" sth ", " south "), (" hts ", " heights "), (" upr ", " upper ")):
        s = s.replace(a, b)
    return " ".join(s.split())


def run(t=TARGET):
    reg = register(t)
    codes = sorted(reg)
    rows = year5_reading(t)
    acc = {}
    for c, yr, y in rows:
        acc.setdefault((c, yr), []).append(y)
    assert {c for c, _ in acc} == set(codes)
    groups, SSW, N = [], 0.0, 0
    for c in codes:
        g = []
        for yr in (2024, 2025, 2026):
            v = acc[(c, yr)]
            m = sum(v) / len(v)
            SSW += sum((x - m) ** 2 for x in v)
            N += len(v)
            g.append((len(v), m))
        groups.append(g)
    tu, tv, s2, mu, iters = reml_em(groups, SSW, N)
    prob = {}
    for c, g in zip(codes, groups):
        w = [1.0 / (tv + s2 / n) for n, _ in g]
        P = 1.0 / tu + sum(w)
        m = mu + sum(wi * (y - mu) for wi, (_, y) in zip(w, g)) / P
        prob[c] = _phi((480.0 - m) / math.sqrt(1.0 / P))
    total = sum(prob.values())
    regions = sorted({v["region"] for v in reg.values()})
    out = {"components": {"tu": tu, "tv": tv, "s2": s2, "mu": mu, "em_iterations": iters},
           "places_exact": total, "places": int(round(total)),
           "shares": {r: 100.0 * sum(prob[c] for c in codes if reg[c]["region"] == r) / total for r in regions}}
    # census: the year level the student system held on the day of the run. At an F-6 school whose rows run
    # to Year 7 with no Foundation the system had already moved to 2027, so its Years 3 to 6 are 2027's; at
    # every other school 2027 Years 3 to 6 are this year's Years 2 to 5
    held = {}
    with open(os.path.join(t, "census_aug2026_enrolments.csv"), newline="") as f:
        for r in csv.DictReader(f):
            c = "%04d" % int(r["school_no"])
            if c in reg:
                held.setdefault(c, {})[r["year_level"]] = int(r["headcount"])
    stu, out["rolled"] = {}, 0
    for c, lv in held.items():
        assert reg[c]["levels"] == "F-6"
        moved = "7" in lv and "F" not in lv
        assert moved or ("F" in lv and "7" not in lv)
        out["rolled"] += moved
        stu[c] = sum(lv.get(str(k), 0) for k in ((3, 4, 5, 6) if moved else (2, 3, 4, 5)))
    out["students_region"] = {r: sum(prob[c] * stu[c] for c in codes if reg[c]["region"] == r) for r in regions}
    out["students_total"] = sum(prob[c] * stu[c] for c in codes)
    # the coordinators' 2027 plan, matched on name and locality
    key = {(_norm_name(v["name"]), _norm_loc(v["loc"])): c for c, v in reg.items()}
    assert len(key) == len(reg)
    hub_of = {}
    wb = load_workbook(os.path.join(t, "hub_visit_plan_2027.xlsx"), read_only=True)
    for ws in wb.worksheets:
        hdr = None
        for row in ws.iter_rows(values_only=True):
            if hdr is None:
                if row and row[0] == "School":
                    hdr = True
                continue
            if not row or row[0] is None:
                continue
            c = key[(_norm_name(row[0]), _norm_loc(row[1]))]
            assert c not in hub_of
            hub_of[c] = ws.title
    assert set(hub_of) == set(codes)
    km = {}
    with open(os.path.join(t, "hub_school_road_distances.csv"), newline="") as f:
        for r in csv.DictReader(f):
            km[(r["hub"], "%04d" % int(r["school_no"]))] = float(r["road_km"])
    hubs = sorted(set(hub_of.values()))
    out["students_hub"] = {h: sum(prob[c] * stu[c] for c in codes if hub_of[c] == h) for h in hubs}
    out["km_hub"] = {h: sum(prob[c] * 2 * km[(h, c)] for c in codes if hub_of[c] == h) for h in hubs}
    out["km_total"] = sum(out["km_hub"].values())
    out["chart_order"] = sorted(hubs, key=lambda h: -out["km_hub"][h])
    return out


if __name__ == "__main__":
    o = run(sys.argv[1] if len(sys.argv) > 1 else TARGET)
    print("places %d (%.4f)" % (o["places"], o["places_exact"]))
    for k in ("shares", "students_region", "students_hub", "km_hub"):
        print(k, {a: round(b, 3) for a, b in o[k].items()})
    print("students total %.3f, km total %.3f, chart order %s" % (o["students_total"], o["km_total"], o["chart_order"]))
