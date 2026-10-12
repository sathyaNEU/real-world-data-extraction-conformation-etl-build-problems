"""The pack's documents: the analysis standard (PDF), the Partnership guidelines (DOCX) and the data team's
notes (TXT)."""
import io
import re
import zipfile

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import postprocess as PP

ORG = "Catholic Schools Office"


def _styles():
    ss = getSampleStyleSheet()
    return dict(
        body=ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=9.8, leading=13.4, spaceAfter=6),
        head=ParagraphStyle("h", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=11, spaceBefore=10,
                            spaceAfter=4),
        title=ParagraphStyle("t", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=15, leading=19),
        small=ParagraphStyle("s", parent=ss["Normal"], fontName="Helvetica", fontSize=8, leading=10.5,
                             textColor=colors.HexColor("#444444")),
    )


def _footer(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(20 * mm, 11 * mm, text)
        canvas.drawRightString(A4[0] - 20 * mm, 11 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


STANDARD = [
    ("h", "1  Purpose and scope"),
    ("p", "1.1  This standard sets out how the Catholic Schools Office analyses NAPLAN results when it describes "
          "the results of its schools, or of groups of its schools, for its own planning. It applies to the "
          "Learning and Teaching directorate, the system data team and any contractor working on the Office's "
          "behalf."),
    ("p", "1.2  It does not govern the reports the test authority or the national body publish, the school "
          "reports sent to families, or the school-level reports principals receive from the results portal. "
          "Those are produced outside the Office and are used as released."),
    ("p", "1.3  The standard covers the 2024, 2025 and 2026 sittings. It will be reviewed after the 2027 results "
          "are released."),
    ("h", "2  Data"),
    ("p", "2.1  Analyses use the NAPLAN results the state test authority releases for each sitting. A student's "
          "result counts in a test when the student participated; absent, exempt and withdrawn students, and "
          "students whose test was abandoned, have no scale score and do not enter an analysis."),
    ("p", "2.2  NAPLAN scale scores are on one measurement scale across the 2024, 2025 and 2026 sittings, so "
          "results from those sittings are analysed together without adjustment."),
    ("p", "2.3  A result belongs to the school the student sat the test at, as the release records it. Students "
          "who change school after a sitting are not re-assigned."),
    ("p", "2.4  Results are held by the system data team in the form each sitting was supplied: the legacy "
          "extract for 2024, the results-portal download for 2025 and the data-platform export for 2026. The data "
          "team's notes describe the fields of each."),
    ("h", "3  Model"),
    ("p", "3.1  School results in a test are analysed with a random-effects model: a fixed overall mean, a random "
          "effect for the school, a random effect for the year group within the school (one year level at one "
          "sitting), and the student residual."),
    ("p", "3.2  Variance components are estimated by restricted maximum likelihood."),
    ("p", "3.3  Every school with results enters the model, whatever the size of its year groups."),
    ("p", "3.4  Calculations for a school use the overall mean and the variance components at their fitted values."),
    ("p", "3.5  Each test and year level is modelled on its own results. Results in other tests or year levels "
          "are not added to the model."),
    ("h", "4  Working practice"),
    ("p", "4.1  Every analysis is run from a script kept in the data team's repository, so that it can be run "
          "again on the same extracts with the same result."),
    ("p", "4.2  Results for an individual school are not circulated beyond the Learning and Teaching directorate, "
          "the school's principal and the regional consultant without the Director's approval."),
    ("p", "4.3  An analysis that departs from section 3 states the departure and the reason for it on its first "
          "page."),
    ("h", "5  Roles"),
    ("p", "5.1  The Director, Learning and Teaching, owns this standard. The system data team lead maintains the "
          "extracts and the analysis scripts and answers questions on the standard."),
]


def standard_pdf(path):
    st = _styles()
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=22 * mm, rightMargin=22 * mm, topMargin=20 * mm,
                            bottomMargin=20 * mm, title="LT-AS-04 Analysing school results in NAPLAN",
                            author=ORG, subject="Analysis standard", creator=ORG, producer=ORG)
    story = [Paragraph("%s &nbsp;|&nbsp; Learning and Teaching" % ORG, st["small"]), Spacer(1, 4),
             Paragraph("LT-AS-04 &nbsp; Analysing school results in NAPLAN", st["title"]), Spacer(1, 8)]
    ctl = Table([["Document", "LT-AS-04", "Version", "3"],
                 ["Approved", "12 February 2026", "Owner", "Director, Learning and Teaching"],
                 ["Supersedes", "LT-AS-04 version 2 (March 2023)", "Contact", "System data team lead"]],
                colWidths=[24 * mm, 58 * mm, 20 * mm, 64 * mm])
    ctl.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 8.4),
                             ("FONT", (0, 0), (0, -1), "Helvetica-Bold", 8.4),
                             ("FONT", (2, 0), (2, -1), "Helvetica-Bold", 8.4),
                             ("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [ctl, Spacer(1, 8)]
    for kind, text in STANDARD:
        story.append(Paragraph(text, st["head"] if kind == "h" else st["body"]))
    story += [Spacer(1, 4), Paragraph("Document history", st["head"])]
    cell = ParagraphStyle("c", parent=st["body"], fontSize=8.4, leading=10.6, spaceAfter=0)
    hist = Table([["Version", "Date", "Change"],
                  ["1", "June 2019", Paragraph("First issue, 2018 and 2019 sittings, paper tests.", cell)],
                  ["2", "March 2023", Paragraph("Online sittings from 2022; results from 2023 on the new scale "
                                                "only.", cell)],
                  ["3", "February 2026", Paragraph("Scope moved to the 2024 to 2026 sittings; 2.4 and 3.4 added; "
                                                   "4.2 brought into line with the Office's privacy policy.", cell)]],
                 colWidths=[18 * mm, 28 * mm, 120 * mm])
    hist.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 8.4),
                              ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.4),
                              ("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [hist, Spacer(1, 12),
              Paragraph("Approved by the Director, Learning and Teaching, 12 February 2026. Printed copies are "
                        "uncontrolled; the current version is the one on the Office intranet.", st["small"])]
    doc.build(story, onFirstPage=_footer("LT-AS-04 v3  |  %s" % ORG), onLaterPages=_footer("LT-AS-04 v3  |  %s" % ORG))
    PP.normalize_pdf(path, b"D:20260212103400+11'00'", ORG)


GUIDE = [
    ("t", "Reading Partnership 2027"),
    ("s", "Programme guidelines  |  Learning and Teaching  |  October 2026"),
    ("h", "1  The Partnership"),
    ("p", "The Reading Partnership places a specialist reading teacher in each partner school for the 2027 school "
          "year. The specialist visits once a fortnight from the hub that services the school and works with the "
          "Year 3 to 6 classes, alongside the class teachers and the school's literacy leader. Specialists are "
          "engaged through the Office's panel provider for specialist literacy services and are based at the "
          "hubs."),
    ("h", "2  Partner schools"),
    ("p", "Every primary school whose underlying Year 5 reading average is below 480 on the NAPLAN scale is a "
          "partner school for 2027. A school's underlying Year 5 reading average is the average its Year 5 "
          "students would reach over many years if the school kept teaching as it did across the 2024 to 2026 "
          "sittings."),
    ("h", "3  Places"),
    ("p", "One place is commissioned for each partner school. Partner schools are confirmed by the literacy "
          "reviews held in every primary school in Term 1 2027; the places are commissioned at the executive's "
          "meeting of 18 November 2026, before the reviews."),
    ("h", "4  Travel"),
    ("p", "Specialists travel by road from their hub to each partner school and back for every visit."),
    ("h", "5  Timeline"),
    ("b", "18 November 2026: executive meeting, places commissioned"),
    ("b", "December 2026: provider contract signed; specialists allocated to hubs"),
    ("b", "Term 1 2027 (weeks 2 to 8): literacy reviews in every primary school"),
    ("b", "Term 1 2027 (week 9): partner schools confirmed to principals"),
    ("b", "Term 2 2027 (week 1): fortnightly visits begin"),
    ("h", "6  Contacts"),
    ("p", "Laura Proctor, Partnership lead, Learning and Teaching. Questions about hub arrangements go to the hub "
          "coordinators; questions about results data go to the system data team."),
]


def guidelines_docx(path):
    from docx import Document
    from docx.shared import Pt
    d = Document()
    stl = d.styles["Normal"]
    stl.font.name = "Calibri"
    stl.font.size = Pt(11)
    for kind, text in GUIDE:
        if kind == "t":
            d.add_heading(text, level=0)
        elif kind == "s":
            p = d.add_paragraph(text)
            p.runs[0].italic = True
        elif kind == "h":
            d.add_heading(text, level=1)
        elif kind == "b":
            d.add_paragraph(text, style="List Bullet")
        else:
            d.add_paragraph(text)
    cp = d.core_properties
    cp.author = "Laura Proctor"
    cp.last_modified_by = "Laura Proctor"
    cp.title = "Reading Partnership 2027 guidelines"
    cp.comments = ""
    cp.keywords = ""
    cp.subject = ""
    cp.category = ""
    cp.revision = 7
    import datetime as dt
    cp.created = dt.datetime(2026, 9, 14, 9, 20)
    cp.modified = dt.datetime(2026, 10, 21, 15, 48)
    buf = io.BytesIO()
    d.save(buf)
    normalize_docx(buf.getvalue(), path)


def normalize_docx(data, path):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        parts = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    drop = {"word/stylesWithEffects.xml", "docProps/thumbnail.jpeg"}
    drop |= {n for n, _ in parts if n.startswith("customXml/")}
    out = []
    for name, b in parts:
        if name in drop:
            continue
        if name.endswith(".rels") or name == "[Content_Types].xml":
            b = re.sub(rb'<Relationship [^>]*Target="[^"]*(stylesWithEffects\.xml|thumbnail\.jpeg|customXml/[^"]*)"/>',
                       b"", b)
            b = re.sub(rb'<Override PartName="/(word/stylesWithEffects\.xml|docProps/thumbnail\.jpeg|customXml/[^"]*)"'
                       rb'[^>]*/>', b"", b)
            b = re.sub(rb'<Default Extension="jpeg"[^>]*/>', b"", b)
        if name.endswith((".xml", ".rels")):
            b = b.replace(b"<?xml version='1.0' encoding='UTF-8' standalone='yes'?>\n", PP.DECL)
        if name == "docProps/app.xml":
            b = re.sub(rb"<Application>[^<]*</Application>", b"<Application>Microsoft Office Word</Application>", b)
            b = re.sub(rb"<AppVersion>[^<]*</AppVersion>", b"<AppVersion>16.0000</AppVersion>", b)
            b = re.sub(rb"<Company>[^<]*</Company>", b"", b)
        if name == "docProps/core.xml":
            b = re.sub(rb"<cp:lastPrinted>[^<]*</cp:lastPrinted>", b"", b)
            b = b.replace(b"python-docx", b"")
        out.append((name, b))
    with zipfile.ZipFile(path, "w") as z:
        for name, b in out:
            zi = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, b)


def notes_txt(path, files):
    lines = [
        "Catholic Schools Office, system data team",
        "Notes on the extracts in this folder",
        "Joshua Nguyen, data team lead, 9 November 2026",
        "",
        "These are the files the data team holds for 2027 planning. Everything here is internal to the Office;",
        "results for individual schools are handled under LT-AS-04 section 4.2.",
        "",
        "NAPLAN RESULTS",
        "",
        "%s" % files["r2024"],
        "  2024 sitting, from the legacy results system before it was retired. One row per student per test.",
        "  EXTRACT_DT  date of the extract (yyyymmdd)",
        "  SCHL_CD     school code, four characters",
        "  STU_ID      student number",
        "  YR_LVL      year level, 03 or 05",
        "  TEST_CD     RD reading, WR writing, SP spelling, GP grammar and punctuation, NU numeracy",
        "  PART_CD     P participated, A absent, E exempt, W withdrawn, S sanctioned abandonment",
        "  SCALE_SCR   NAPLAN scale score, one decimal; blank unless PART_CD is P",
        "",
        "%s" % files["r2025"],
        "  2025 sitting, downloaded from the results portal on 22 August 2025. A sheet for each year level, a",
        "  title block above the header row. Scale scores are whole numbers in the portal download.",
        "",
        "%s" % files["r2026"],
        "  2026 sitting, exported from the data platform on 28 August 2026.",
        "  school_id             school code, four characters",
        "  student_ref           student number",
        "  year_level            Year 3 or Year 5",
        "  domain                reading, writing, spelling, grammar and punctuation, numeracy",
        "  participation_status  participated, absent, exempt, withdrawn, sanctioned abandonment",
        "  scaled_score          NAPLAN scale score, whole number; blank unless participated",
        "  test_window           month of the online test window",
        "",
        "  All three cover every student enrolled in Years 3 and 5 at our primary schools at the time of",
        "  testing. Student numbers in all three are the data team's study numbers, issued in place of the",
        "  state student number; a student keeps the same one between sittings.",
        "",
        "SCHOOLS AND ENROLMENTS",
        "",
        "%s" % files["register"],
        "  The school register as refreshed on 3 March 2026: every school the Office runs, primary and",
        "  secondary, with its code, locality and region. No school has opened, closed or merged since, and",
        "  none is planned for 2027.",
        "",
        "%s" % files["census"],
        "  Headcount by year level of the students on each school's roll on census day (7 August 2026),",
        "  every school, run from the student system on 4 November 2026. Year level is as the student",
        "  system held it on the day of the run. Students move up one year level at the start of each school",
        "  year. The run writes the school number as a number, so codes that start with zero have lost it.",
        "",
        "HUBS AND TRAVEL",
        "",
        "%s" % files["directory"],
        "  Facilities' directory of sites and the hub that services each, from its annual February refresh.",
        "  Kept by Facilities (Alice Wood).",
        "",
        "%s" % files["plan"],
        "  The hub coordinators' servicing plan for 2027, a sheet for each hub, schools listed by name and",
        "  locality as the coordinators keep them. Version of 6 November 2026 (Diana Hardy compiles it).",
        "",
        "%s" % files["road"],
        "  One-way road distance and drive time from the routing service, refreshed 2 November 2026. Schools",
        "  in the two metropolitan regions were run against the two metropolitan hubs; schools in the other",
        "  regions against all five. School numbers as in the census.",
        "",
        "OTHER",
        "",
        "%s" % files["rates"],
        "  The panel provider's 2027 schedule of rates for specialist literacy services.",
        "",
        "%s" % files["pl"],
        "  Professional learning register for literacy courses, 2026 school year to date, one row per",
        "  enrolment. School codes as in the register; staff numbers are study numbers too.",
        "",
        "  All files are Office records, extracted by the data team, apart from the provider's schedule of",
        "  rates. NAPLAN results are released to the Office by the state test authority for its own schools.",
        "",
    ]
    with open(path, "w", newline="") as f:
        f.write("\r\n".join(lines))
