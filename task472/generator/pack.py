"""The shipped extracts, each in the layout and clock of the system that writes it: the Authority's terminal operating
system (crane hours and visits, Eastern time), the SRT workbook (UTC), the pilot station's sailings (Eastern
time), the Salt Point tide predictions (local standard time), the vessel and crane registers, the Authority's
maintenance system (orders, status changes and labor, Eastern time), the crane makers' service timesheets, the
SRT maintenance log (UTC), the AIS berth segments (UTC), the yard equipment orders and the 2027 berth windows."""
import csv
import datetime as dt

import numpy as np

import world as W

H, MIN, DAY = W.H, W.MIN, W.DAY
FILES = dict(
    tos="tos_crane_activity_2026.csv",
    visits="tos_vessel_visits_2026.csv",
    rvt="srt_tos_export_fy2026.xlsx",
    pilot="pilot_station_sailings_fy2026.csv",
    tides="salt_point_tide_predictions_2025_2026.txt",
    vessels="vessel_register_extract_2026-10-05.json",
    register="sts_crane_register.csv",
    wo="cmms_sts_work_orders_fy2026.csv",
    status="cmms_asset_status_history_fy2026.csv",
    labor="cmms_labor_transactions_fy2026.csv",
    oem="crane_maker_service_timesheets_fy2026.xlsx",
    jobs="srt_maintenance_jobs_2025-26.xlsx",
    ais="ais_berth_segments_fy2026.csv",
    yard="cmms_yard_equipment_work_orders_fy2026.csv",
    windows="berth_windows_fy2027.csv",
    rules="crane_replacement_program_rules.pdf",
    directions="harbor_masters_general_directions_2025.pdf",
    dictionary="data_dictionary.txt",
    manifest="export_manifest.csv",
)
TOS_FROM, TOS_TO = W.ts(2025, 6, 23), W.ts(2026, 7, 7)
PILOT_FROM, PILOT_TO = W.ts(2025, 6, 20), W.ts(2026, 7, 10)
AIS_FROM, AIS_TO = W.ts(2025, 6, 20, 0, 0, W.UTC), W.ts(2026, 7, 10, 0, 0, W.UTC)
WO_FROM, WO_TO = W.ts(2025, 6, 1), W.ts(2026, 10, 1)
PULL = W.ts(2026, 10, 5, 6, 10)
LINE = {"Corvane Shipping": "CRV", "Eastmere Container Line": "EML", "Hallvik Line": "HAL",
        "Pacific Arrow Lines": "PAL", "Tessaline": "TSL", "Morrow Kade Container": "MKC", "Brightwater Liner": "BWL",
        "Ostrel Marine": "OST"}
USERS = dict(ORD=["KWILLIAMS", "ORD-OPS2", "ORD-MNT1"], PNC=["PNC-MNT1", "PNC-OPS1", "DGILL"],
             GNP=["GNP-MNT1", "GNP-OPS1"], HLY=["HLY-MNT2", "HLY-OPS1", "DGILL"], SLT=["JWOOD", "SLT-MNT1"],
             MRW=["JWOOD", "MRW-MNT1"])
PILOTS = ["RTH", "MKO", "JDL", "SPV", "AGR", "TWB", "LNC", "DFR", "CMK", "EJS", "PBA", "KRY"]


def lstr(t, sec=True):
    return W.local(t).strftime("%Y-%m-%d %H:%M:%S" if sec else "%Y-%m-%d %H:%M")


def ustr(t, sec=True):
    return W.utc(t).strftime("%Y-%m-%d %H:%M:%S" if sec else "%Y-%m-%d %H:%M")


def unaive(t):
    return W.utc(t).replace(tzinfo=None)


def write_csv(path, cols, rows, crlf=False):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh, lineterminator="\r\n" if crlf else "\n")
        wr.writerow(cols)
        n = 0
        for r in rows:
            wr.writerow(r)
            n += 1
    return n


# ---------------------------------------------------------------------------------------------- crane hours
def run_hours(r, rng):
    """One crane's run as hourly rows: (utc hour start, moves, first move, last move, lowest bay, highest bay)."""
    t = r.to_clock(r.moves_w)
    n = len(t)
    k = len(r.bays)
    cuts = np.cumsum(rng.dirichlet(np.ones(k) * 4.0))
    pos = np.arange(n) / max(n, 1)
    bay_idx = np.searchsorted(cuts, pos, side="right").clip(0, k - 1)
    bays = np.array(r.bays)[bay_idx]
    hour = (t // H) * H
    out = []
    for h in np.unique(hour):
        m = hour == h
        out.append((int(h), int(m.sum()), int(t[m].min()), int(t[m].max()), int(bays[m].min()), int(bays[m].max())))
    return out


def tos_rows(x):
    """Every crane-hour of every call, in both systems' grains; the hour bucket is the system's own clock hour."""
    rng = np.random.default_rng(W.SEED + 401)
    auth, rvt = [], []
    for c in x["calls"]:
        harbor = W.TERM[c.term]["water"] == "harbor"
        for r in c.runs:
            for (h, mv, f, l, b0, b1) in run_hours(r, rng):
                if not (TOS_FROM <= h < TOS_TO):
                    continue
                hc = int(rng.integers(1, 5)) if rng.random() < 0.14 else 0
                if harbor:
                    auth.append((h, c.term, c.visit, r.crane, mv, hc, f, l, b0, b1))
                else:
                    rvt.append((h, c.term, c.visit, W.CR[r.crane]["tag"], mv, f, l, b0, b1))
    auth.sort(key=lambda z: (z[0], z[1], z[3]))
    rvt.sort(key=lambda z: (z[1], z[0], z[3]))
    return auth, rvt


TOS_COLS = ["TERMINAL", "VISIT_ID", "CRANE_ID", "HOUR_START", "MOVES", "HATCH_COVERS", "FIRST_MOVE", "LAST_MOVE",
            "BAY_FROM", "BAY_TO"]


def tos_csv_rows(auth):
    for h, term, visit, crane, mv, hc, f, l, b0, b1 in auth:
        yield [term, visit, crane, W.local(h).strftime("%Y-%m-%d %H:00"), mv, hc, lstr(f), lstr(l),
               "%02d" % b0, "%02d" % b1]


VISIT_COLS = ["VISIT_ID", "TERMINAL", "BERTH", "VESSEL_NAME", "IMO_NO", "LINE_OP", "SERVICE", "IB_VOYAGE",
              "OB_VOYAGE", "ETA", "PHASE"]


def voyage(rng):
    n = int(rng.integers(100, 999))
    return "%03dW" % n, "%03dE" % n


def visit_rows(x):
    rng = np.random.default_rng(W.SEED + 402)
    out = []
    for c in x["calls"]:
        if W.TERM[c.term]["water"] != "harbor":
            continue
        if not (TOS_FROM - 3 * DAY <= c.first <= TOS_TO):
            continue
        ib, ob = voyage(rng)
        eta = c.first - c.arr_lag + int(rng.normal(0, 5) * H)
        svc = W.SERVICES[c.term][int(rng.integers(len(W.SERVICES[c.term])))] if c.large else "FDR"
        out.append([c.visit, c.term, c.berth.split()[-1], W.name_on(c.vessel, c.first), str(c.vessel["imo"]),
                    LINE[c.vessel["carrier"]], svc, ib, ob, lstr((eta // MIN) * MIN, sec=False), "CLOSED"])
    out.sort(key=lambda r: r[0])
    return out


def rvt_calls(x):
    rng = np.random.default_rng(W.SEED + 403)
    out = []
    for c in x["calls"]:
        if W.TERM[c.term]["water"] != "river":
            continue
        if not (TOS_FROM - 3 * DAY <= c.first <= TOS_TO):
            continue
        ib, ob = voyage(rng)
        eta = c.first - c.arr_lag + int(rng.normal(0, 5) * H)
        svc = W.SERVICES[c.term][int(rng.integers(len(W.SERVICES[c.term])))] if c.large else "FEEDER"
        out.append([c.visit, c.berth, W.name_on(c.vessel, c.first), int(c.vessel["imo"]),
                    W.call_sign_on(c.vessel, c.first), c.vessel["carrier"], svc, ib + "/" + ob,
                    unaive((eta // MIN) * MIN)])
    out.sort(key=lambda r: r[0])
    return out


# ---------------------------------------------------------------------------------------------- the pilot station
PILOT_COLS = ["JOB_NO", "MOVE", "VESSEL", "CALL_SIGN", "FROM_BERTH", "TO", "NOTICE_RECEIVED", "SAILED", "DRAFT_FWD",
              "DRAFT_AFT", "PILOT", "TUGS"]


def typed(name, rng):
    """A ship's name as a dispatcher typed it."""
    u = rng.random()
    if u < 0.05:
        return "M/V " + name
    if u < 0.08:
        return name.replace(" ", "  ", 1)
    if u < 0.10:
        return name.title()
    return name


def pilot_rows(x):
    rng = np.random.default_rng(W.SEED + 409)
    calls = sorted([c for c in x["calls"] if PILOT_FROM <= c.sail < PILOT_TO], key=lambda c: (c.sail, c.term))
    out = []
    n = {}
    for c in calls:
        yy = W.local(c.sail).year % 100
        n[yy] = n.get(yy, 1180 if yy == 25 else 0) + int(rng.integers(1, 4))
        aft = c.draft
        fwd = max(aft - float(rng.uniform(0.0, 1.6)), aft - 1.6)
        fa, ia = W.ft_in(aft)
        ff, if_ = W.ft_in(fwd)
        if (ff, if_) > (fa, ia):
            ff, if_ = fa, ia
        big = c.vessel["loa"] / W.FT >= W.LONG_FT
        tugs = (3 if big else 2) + (1 if W.TERM[c.term]["water"] == "river" and c.large else 0)
        out.append(["%02d-%05d" % (yy, n[yy]), "OUT", typed(W.name_on(c.vessel, c.sail), rng),
                    W.call_sign_on(c.vessel, c.sail), W.PILOT_BERTH[c.berth], "SEA",
                    lstr(c.notice, False), lstr(c.sail, False), "%d-%02d" % (ff, if_), "%d-%02d" % (fa, ia),
                    PILOTS[int(rng.integers(len(PILOTS)))], tugs])
    return out


# ---------------------------------------------------------------------------------------------- tides
def tides_txt(x):
    """Salt Point predictions in the layout of a published annual table, local standard time all year."""
    lines = ["Ternmouth Port Authority - Hydrographic Services",
             "Annual Tide Predictions",
             "Station: SALT POINT, SALT RIVER ENTRANCE",
             "Source: harmonic predictions, Ternmouth Port Authority",
             "Time Zone: LST (UTC-05:00); add one hour when daylight saving time is in effect",
             "Datum: MLLW (heights in feet)",
             "Interval Type: High/Low Tide Predictions",
             "From 2025/06/01 To 2026/07/31",
             "",
             "Date\tDay\tTime\tPred(Ft)\tPred(cm)\tHigh/Low"]
    lo, hi = W.ts(2025, 6, 1, 0, 0, W.LST), W.ts(2026, 8, 1, 0, 0, W.LST)
    for t, k, h in x["tides"]:
        if not (lo <= t < hi):
            continue
        d = dt.datetime.fromtimestamp(t, W.LST)
        lines.append("%s\t%s\t%s\t%.1f\t%d\t%s" % (d.strftime("%Y/%m/%d"), d.strftime("%a"),
                                                   d.strftime("%I:%M %p"), h, int(round(h * 30.48)), k))
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------------------------------------- registers
REGISTRY = {"Liberia": "Monrovia", "Marshall Islands": "Majuro", "Panama": "Panama", "Singapore": "Singapore",
            "Malta": "Valletta", "Hong Kong": "Hong Kong", "Portugal": "Madeira", "Cyprus": "Limassol",
            "Denmark": "Copenhagen", "Greece": "Piraeus", "Germany": "Hamburg"}


def vessel_json(x):
    vs = []
    for v in sorted(x["fleet"], key=lambda v: v["imo"]):
        vs.append(dict(imo=v["imo"], name=v["name"],
                       former_names=[dict(name=f["name"], until=f["until"]) for f in v.get("former", [])],
                       call_sign=v["call_sign"], mmsi=v["mmsi"], port_of_registry=REGISTRY[v["flag"]], teu_nominal=v["teu"],
                       loa_m=v["loa"], beam_m=v["beam"], summer_draft_m=v["draft"], dwt=v["dwt"], built=v["built"],
                       operator=v["carrier"]))
    return dict(extract=dict(source="Ternmouth Port Authority vessel register", pulled="2026-10-05T06:10:00-04:00",
                             records=len(vs), note="current particulars; former names with the last date each was"
                             " carried"), vessels=vs)


REG_COLS = ["CRANE_ID", "TERMINAL", "ASSETNUM", "SRT_TAG", "MAKE", "MODEL", "COMMISSIONED", "OUTREACH_ROWS",
            "SWL_T"]


def register_rows():
    out = []
    swl = {"long": 65, "short": 50, "even": 61}
    for c in W.CRANES:
        spec = W.CR[c[0]]
        out.append([c[0], W.term_now(c[0]), spec["asset"], spec["tag"] or "", spec["make"], spec["model"],
                    spec["year"], spec["outreach"], 65 if spec["make"] == "ZPMC" else (40 if spec["year"] < 1997
                                                                                      else swl[spec["role"]])])
    out.sort(key=lambda r: (["ORD", "PNC", "GNP", "HLY", "SLT", "MRW"].index(r[1]), r[0]))
    return out


# ---------------------------------------------------------------------------------------------- maintenance
PM_TEXT = ["Monthly PM, STS crane", "Monthly PM incl. spreader service", "Monthly PM, rope inspection and lube",
           "Monthly PM, brakes and limit switches"]
INSP_TEXT = ["Annual load test and certification", "Structural inspection, boom and girder", "Wire rope NDT inspection"]
CM_TEXT = ["Replace cab seat", "Repair walkway grating", "Replace floodlight", "Repaint boom tip markings",
           "Replace cab wiper motor", "Fix lift intercom", "Replace handrail section", "Recalibrate wind anemometer"]
DMG_TEXT = ["Spreader struck hatch coaming, flipper arm bent - ops incident report filed",
            "Trolley festoon struck by vessel lashing bridge - damage claim to line",
            "Gantry bogie struck by straddle carrier - yard incident",
            "Boom tip light and cable struck by ship's crane - damage claim to line",
            "Spreader headblock damaged on landing, operator incident report"]
SUFFIX = ["", "", "", " - reset and tested", " - parts replaced", " - adjusted", " - operator reported",
          " - found on inspection", " - temporary repair", " - see attached notes"]


def authority_events(x):
    return [e for e in x["maint"] if e["system"] in ("authority", "both") and e["kind"] != "MOVE"]


def legacy_events(x):
    return [e for e in x["maint"] if e["system"] in ("legacy", "both") and e["kind"] != "MOVE"]


def yard_orders(rng):
    """The yard equipment orders: same system, same sites, other asset classes."""
    kinds = [("RTG", 26, 1 / 900), ("STRADDLE", 34, 1 / 700), ("REACH", 18, 1 / 800), ("TT", 60, 1 / 1100)]
    out = []
    for kind, n, lam in kinds:
        for i in range(n):
            asset = "4%02d%03d" % ({"RTG": 1, "STRADDLE": 2, "REACH": 3, "TT": 4}[kind], i + 11)
            term = W.CODES[i % 6]
            t = W.ts(2025, 6, 1) + int(rng.uniform(0, 5) * DAY)
            while t < W.ts(2026, 9, 28):
                t += int(rng.exponential(1 / lam) * H)
                d = W.mins(rng.uniform(0.4, 6.0))
                out.append(dict(asset=asset, term=term, kind="BD", start=(t // MIN) * MIN, dur=d, cls=kind))
            t = W.ts(2025, 6, 1) + int(rng.uniform(0, 20) * DAY)
            while t < W.ts(2026, 9, 28):
                d = W.mins(rng.uniform(2.0, 6.0))
                out.append(dict(asset=asset, term=term, kind="PM", start=(t // MIN) * MIN, dur=d, cls=kind))
                t += int(rng.uniform(40, 62) * DAY)
    return out


WO_COLS = ["WONUM", "ASSETNUM", "SITEID", "WORKTYPE", "WOPRIORITY", "STATUS", "PROBLEMCODE", "DESCRIPTION",
           "REPORTDATE", "OUT_OF_SERVICE", "RETURNED_TO_SERVICE", "ACTSTART", "ACTFINISH", "CLOSEDATE", "REPORTEDBY"]
STATUS_COLS = ["ASSETNUM", "ISRUNNING", "CHANGEDATE", "CODE", "WONUM", "CHANGEBY"]


def order_book(x):
    """Number every order the Authority's system holds (cranes and yard equipment share the sequence) and lay out the
    crane orders, their status changes (with the downtime of each crane's move) and the yard orders."""
    rng = np.random.default_rng(W.SEED + 404)
    yard = yard_orders(np.random.default_rng(W.SEED + 405))
    ev = [e for e in authority_events(x) if WO_FROM <= e["periods"][0][0] < WO_TO]
    items = [(e["periods"][0][0] + e["reported_lag"], 0, e) for e in ev]
    for i, y in enumerate(yard):
        y["seq"] = i
    items += [(y["start"] + 3 * MIN, 1, y) for y in yard if WO_FROM <= y["start"] < WO_TO]
    items.sort(key=lambda z: (z[0], z[1], z[2]["seq"]))
    num = 1584217
    trng = np.random.default_rng(W.SEED + 414)
    pick = lambda xs: xs[int(trng.integers(len(xs)))]
    sts, status, yrows = [], [], []
    for rep, kind, e in items:
        num += int(rng.integers(1, 4))
        wonum = str(num)
        if kind == 1:
            s, d = e["start"], e["dur"]
            closed = s + d + int(rng.uniform(1, 9) * DAY)
            e["wonum"] = wonum
            yrows.append([wonum, e["asset"], "TMH", e["kind"], 1 if e["kind"] == "BD" else 3,
                          "CLOSE" if closed < PULL else "COMP", e["cls"] + ("-BD" if e["kind"] == "BD" else "-PM"),
                          ("%s breakdown" % e["cls"].title()) if e["kind"] == "BD" else "Scheduled PM",
                          lstr(rep, False), lstr(s, False), lstr(s + d, False),
                          lstr(s + 12 * MIN, False), lstr(s + d - 6 * MIN, False),
                          (lstr(closed, False) if closed < PULL else ""), USERS[e["term"]][0]])
            continue
        per = e["periods"]
        s = per[0][0]
        asset = W.CR[e["crane"]]["asset"]
        users = USERS[e["term"]]
        by = users[int(rng.integers(len(users)))]
        e["wonum"] = wonum
        if e["kind"] == "CM":
            a = s + int(rng.uniform(1, 30) * H)
            fin = a + int(rng.uniform(0.5, 4) * H)
            closed = fin + e["closed_lag"]
            sts.append([wonum, asset, "TMH", "CM", 3, "CLOSE" if closed < PULL else "COMP", "GENERAL",
                        pick(CM_TEXT), lstr(s, False), "", "", lstr(a, False),
                        lstr(fin, False), (lstr(closed, False) if closed < PULL else ""), by])
            continue
        last = per[-1][1]
        la, lb = e["labour_lag"]
        a = s + la
        fin = max(last - lb, a + 10 * MIN)
        closed = last + e["closed_lag"]
        if e["kind"] == "BD" and e.get("damage"):
            wt, pri, code = "DMG", 1, "DAMAGE"
            text = pick(DMG_TEXT)
        elif e["kind"] == "BD":
            wt, pri, code = "BD", 1, e["fault"]
            text = e["text"] + pick(SUFFIX)
        elif e["kind"] == "PM":
            wt, pri, code, text = "PM", 3, "PM-MONTHLY", pick(PM_TEXT)
        else:
            wt, pri, code = "INSP", 2, "INSPECT"
            text = INSP_TEXT[0] if e["src"] == "annual" else INSP_TEXT[1 + int(trng.integers(2))]
        sts.append([wonum, asset, "TMH", wt, pri, "CLOSE" if closed < PULL else "COMP", code, text,
                    lstr(rep, False), lstr(s, False), lstr(last, False), lstr(a, False), lstr(fin, False),
                    (lstr(closed, False) if closed < PULL else ""), by])
        stat_code = {"BD": "BRKDWN", "PM": "PLANNED", "INSP": "PLANNED"}[e["kind"]]
        for p0, p1 in per:
            status.append([asset, 0, lstr(p0, False), stat_code, wonum, by, p0])
            status.append([asset, 1, lstr(p1, False), stat_code, wonum, by, p1])
    # the move crew books the crane out when it is made safe for the move and back once it is load tested on its new rail
    prng = np.random.default_rng(W.SEED + 413)
    for e in x["maint"]:
        if e["kind"] == "MOVE":
            a, b = e["periods"][0]
            a -= int(prng.uniform(25, 95)) * MIN
            b += int(prng.uniform(70, 330)) * MIN
            asset = W.CR[e["crane"]]["asset"]
            if W.CR[e["crane"]]["tag"] and a < W.CUTOVER[0]:
                continue
            status.append([asset, 0, lstr(a, False), "PROJECT", "", "DGILL", a])
            status.append([asset, 1, lstr(b, False), "PROJECT", "", "DGILL", b])
    status.sort(key=lambda r: (r[-1], r[0], r[1]))
    status = [r[:-1] for r in status]
    return sts, status, yrows


LABOR_COLS = ["LABTRANSID", "WONUM", "LABORCODE", "CRAFT", "TRANSDATE", "REGULARHRS", "PREMIUMPAYCODE", "ENTERDATE",
              "ENTERBY"]


def labor_rows(x, yard):
    """The labor ledger for every order in the Authority's system (cranes and yard equipment): one row per person,
    order and work day; contractors' hours under the code CONTRACT."""
    rng = np.random.default_rng(W.SEED + 410)
    rows = []
    for r in x["labour"]:
        e = r["ev"]
        if e["system"] not in ("authority", "both") or "wonum" not in e:
            continue
        rows.append((r["date"], e["wonum"], r["code"], r["craft"], r["hours"], r["ot"], e["term"]))
    for y in yard:
        if "wonum" not in y:
            continue
        d = W.local(y["start"]).date()
        techs = W.TECHS[y["term"]]
        crew = rng.permutation(len(techs))[:1 if y["kind"] == "PM" else int(rng.integers(1, 3))]
        for k in crew:
            h = max(0.5, round(y["dur"] / H * rng.uniform(0.7, 1.0) * 4) / 4)
            code, craft = techs[int(k)]
            rows.append((d, y["wonum"], code, craft, h, "", y["term"]))
    rows.sort(key=lambda z: (z[0], int(z[1]), z[2]))
    out = []
    lid = 7710342
    for d, wonum, code, craft, h, ot, term in rows:
        lid += int(rng.integers(1, 6))
        enter = dt.datetime.combine(d, dt.time(15, 0)) + dt.timedelta(days=int(rng.integers(0, 3)),
                                                                       minutes=int(rng.integers(0, 240)))
        out.append([lid, wonum, code, craft, d.isoformat(), "%.2f" % h, ot, enter.strftime("%Y-%m-%d %H:%M"),
                    USERS[term][int(rng.integers(len(USERS[term])))]])
    return out


PREFIX = {"ZPMC": "ZPM", "Liebherr": "LCC", "Konecranes": "KCS", "Kalmar": "KFS", "Mitsubishi": "MHI",
          "Atlantic Portainer Services": "APS"}
OEM_COLS = ["SERVICE_CO", "SERVICE_REPORT", "CUSTOMER_REF", "CRANE", "SITE", "ENGINEER", "WORK_DATE", "HOURS"]


def oem_rows(x):
    rng = np.random.default_rng(W.SEED + 411)
    out = []
    n = {}
    for o in sorted(x["oem"], key=lambda o: (o["date"], o["maker"], o["engineer"])):
        e = o["ev"]
        if e["system"] == "legacy":
            ref = e.get("job_no")
            if ref is None:
                continue
        else:
            if "wonum" not in e:
                continue
            ref = ["WO %s", "%s", "WO#%s", "WO-%s"][int(rng.integers(4))] % e["wonum"]
        key = (o["maker"], e["seq"])
        if key not in n:
            n[key] = "%s-%05d" % (PREFIX[o["maker"]], 31000 + int(rng.integers(0, 9000)))
        crane = ["%s", "QC %s", "STS %s", "Crane %s"][int(rng.integers(4))] % e["crane"]
        out.append([o["maker"], n[key], ref, crane, "Ternmouth Port Authority, %s Terminal" % W.NAME[e["term"]],
                    o["engineer"], o["date"].isoformat(), "%.2f" % o["hours"]])
    return out


JOB_COLS = ["JOB_NO", "EQUIP_TAG", "JOB_CLASS", "PRIORITY", "FAULT", "RAISED_UTC", "DOWN_UTC", "UP_UTC", "LABOUR_HRS",
            "STATUS", "RAISED_BY"]
LEGACY_FAULT = {"SPRDR": "Spreader", "HOIST": "Hoist", "TROL": "Trolley", "GANT": "Gantry", "ELEC": "Electrical",
                "BOOM": "Boom", "CAB": "Cab/controls"}


def job_rows(x):
    rng = np.random.default_rng(W.SEED + 406)
    trng = np.random.default_rng(W.SEED + 415)
    pick = lambda xs: xs[int(trng.integers(len(xs)))]
    hours = {}
    for r in x["labour"]:
        if r["kind"] in ("own", "contractor"):
            hours[id(r["ev"])] = hours.get(id(r["ev"]), 0.0) + r["hours"]
    ev = [e for e in legacy_events(x) if e["periods"][0][0] < W.CUTOVER[1]]
    ev.sort(key=lambda e: (e["periods"][0][0], e["seq"]))
    out, n = [], 4105
    for e in ev:
        n += int(rng.integers(1, 3))
        s = e["periods"][0][0]
        tag = W.CR[e["crane"]]["tag"]
        yy = W.utc(s).year % 100
        e["job_no"] = "J%02d-%05d" % (yy, n)
        lab = round(hours.get(id(e), 0.0), 2)
        raised = unaive(((s + e["reported_lag"]) // MIN) * MIN)
        if e["kind"] == "CM":
            out.append([e["job_no"], tag, "DEFECT", "P3", pick(CM_TEXT), raised, None, None, lab,
                        "CLOSED", "MRW-PLANNER" if e["term"] == "MRW" else "SLT-PLANNER"])
            continue
        last = e["periods"][-1][1]
        cls = {"BD": "BRKDN", "PM": "PLND", "INSP": "INSP"}[e["kind"]]
        pri = {"BD": "P1", "PM": "P3", "INSP": "P2"}[e["kind"]]
        fault = (LEGACY_FAULT[e["fault"]] + ": " + e["text"].lower()) if e["kind"] == "BD" else (
            "Monthly service" if e["kind"] == "PM" else "Annual inspection")
        out.append([e["job_no"], tag, cls, pri, fault, raised, unaive(s // MIN * MIN), unaive(last // MIN * MIN), lab,
                    "CLOSED",
                    ["JWOOD", "MRW-SHIFT-A", "MRW-SHIFT-B"][e["seq"] % 3] if e["term"] == "MRW" else
                    ["JWOOD", "SLT-SHIFT-A", "SLT-SHIFT-B"][e["seq"] % 3]])
    return out


# ---------------------------------------------------------------------------------------------- AIS and windows
AIS_COLS = ["SEG_ID", "MMSI", "SHIPNAME", "CALLSIGN", "ZONE", "ZONE_TYPE", "START_UTC", "END_UTC", "N_REPORTS"]


def ais_rows(x):
    out = []
    segs = [s for s in x["ais"] if AIS_FROM <= s["start"] < AIS_TO]
    for i, s in enumerate(segs):
        v = s["vessel"]
        ztype = "ANCHORAGE" if "ANCH" in s["zone"] else "BERTH"
        out.append(["S%07d" % (3104411 + i * 3), s["mmsi"], W.name_on(v, s["start"]), W.call_sign_on(v, s["start"]),
                    s["zone"], ztype, ustr(s["start"]), ustr(s["end"]), s["n"]])
    return out


WINDOW_COLS = ["SERVICE", "CARRIER", "TERMINAL", "BERTH", "ARRIVAL_DAY", "ARRIVAL_TIME", "WINDOW_HOURS", "MAX_TEU",
               "ROTATION_WEEKS", "EFFECTIVE_FROM"]


def window_rows():
    rng = np.random.default_rng(W.SEED + 407)
    days = ["MON", "TUE", "WED", "THU", "FRI", "SAT", "SUN"]
    out = []
    for code, svcs in W.SERVICES.items():
        for s in svcs + ["%sF%d" % (code[0], k) for k in (1, 2)]:
            big = not s.startswith(code[0] + "F")
            carrier = W.CARRIERS[int(rng.integers(len(W.CARRIERS)))][0]
            out.append([s, carrier, code, W.BERTHS[code][int(rng.integers(len(W.BERTHS[code])))].split()[-1],
                        days[int(rng.integers(7))], "%02d:00" % int(rng.integers(0, 24)),
                        int(rng.integers(20, 34)) if big else int(rng.integers(8, 16)),
                        int(rng.choice([13000, 14000, 15000, 16000, 20000, 24000])) if big else
                        int(rng.choice([2800, 4400, 6600, 8800])), int(rng.integers(6, 12)),
                        ["2026-07-01", "2026-07-01", "2026-08-15", "2026-10-01", "2027-01-05"][int(rng.integers(5))]])
    out.sort(key=lambda r: (r[2], r[0]))
    return out


# ---------------------------------------------------------------------------------------------- workbooks
def write_xlsx(path, sheets, widths=None):
    """sheets: list of (title, columns, rows, datetime column indexes)."""
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, cols, rows, dcols in sheets:
        ws = wb.create_sheet(title)
        ws.append(cols)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append(list(r))
        for j in dcols:
            for row in ws.iter_rows(min_row=2, min_col=j + 1, max_col=j + 1):
                for cell in row:
                    if cell.value is not None:
                        cell.number_format = "yyyy-mm-dd hh:mm" if isinstance(cell.value, dt.datetime) else "yyyy-mm-dd"
        for k, col in enumerate(cols):
            ws.column_dimensions[openpyxl.utils.get_column_letter(k + 1)].width = max(10, min(36, len(col) + 4))
        ws.freeze_panes = "A2"
    wb.save(path)
