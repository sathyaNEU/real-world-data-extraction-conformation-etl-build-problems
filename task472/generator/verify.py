"""The independent verifier: recompute every graded figure from the shipped bundle alone, by its own code path (it
imports nothing from the generator), and compare with the golden workbook and the calls table.

    python3 task472/generator/verify.py [target] [golden]
"""
import csv
import datetime as dt
import json
import os
import re
import sys
from collections import defaultdict
from zoneinfo import ZoneInfo

import openpyxl

ET = ZoneInfo("America/New_York")
UTC = dt.timezone.utc
STD = dt.timezone(dt.timedelta(hours=-5))
TERMS = {"ORD": "Ore Dock", "PNC": "Pine Neck", "GNP": "Gannet Point", "HLY": "Halyard", "SLT": "Saltworks",
         "MRW": "Millrace"}
BERTHS = {"ORE DOCK": "ORD", "PINE NECK": "PNC", "GANNET PT": "GNP", "HALYARD": "HLY", "SALTWORKS": "SLT",
          "MILLRACE": "MRW"}


def read(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def sheet(path, name):
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = list(wb[name].iter_rows(values_only=True))
    wb.close()
    return [dict(zip(rows[0], r)) for r in rows[1:]]


def eastern(s, fold=0):
    fmt = "%Y-%m-%d %H:%M:%S" if s.count(":") == 2 else "%Y-%m-%d %H:%M"
    return dt.datetime.strptime(s, fmt).replace(tzinfo=ET, fold=fold).timestamp()


def first_open(t, wins):
    for a, b in wins:
        if b >= t:
            return max(t, a)
    raise ValueError(t)


def main(target, golden):
    f = lambda n: os.path.join(target, n)
    fy0 = dt.datetime(2025, 7, 1, tzinfo=ET).timestamp()
    fy1 = dt.datetime(2026, 7, 1, tzinfo=ET).timestamp()
    # calls: first and last container move by visit, from both terminal systems
    span = defaultdict(lambda: [1e20, -1e20])
    seen = defaultdict(int)
    where = defaultdict(list)
    for r in read(f("tos_crane_activity_2026.csv")):
        k = (r["VISIT_ID"], r["CRANE_ID"], r["HOUR_START"])
        fold = 1 if seen[k] else 0
        seen[k] += 1
        a, b = eastern(r["FIRST_MOVE"], fold), eastern(r["LAST_MOVE"], fold)
        s = span[(r["TERMINAL"], r["VISIT_ID"])]
        s[0], s[1] = min(s[0], a), max(s[1], b)
        where[r["CRANE_ID"]].append((a, r["TERMINAL"]))
    imo = {(r["TERMINAL"], r["VISIT_ID"]): int(r["IMO_NO"]) for r in read(f("tos_vessel_visits_2026.csv"))}
    reg = read(f("sts_crane_register.csv"))
    by_tag = {r["SRT_TAG"]: r["CRANE_ID"] for r in reg if r["SRT_TAG"]}
    by_asset = {r["ASSETNUM"]: r["CRANE_ID"] for r in reg}
    rv = f("srt_tos_export_fy2026.xlsx")
    for r in sheet(rv, "Calls"):
        imo[(r["CALL_REF"][:3], r["CALL_REF"])] = int(r["LLOYDS"])
    for r in sheet(rv, "QC Hours"):
        a = r["FIRST_LIFT_UTC"].replace(tzinfo=UTC).timestamp()
        b = r["LAST_LIFT_UTC"].replace(tzinfo=UTC).timestamp()
        s = span[(r["CALL_REF"][:3], r["CALL_REF"])]
        s[0], s[1] = min(s[0], a), max(s[1], b)
        where[by_tag[r["QC"]]].append((a, r["CALL_REF"][:3]))
    for k in where:
        where[k].sort()
    ships = {v["imo"]: v for v in json.load(open(f("vessel_register_extract_2026-10-05.json")))["vessels"]}
    calls = []
    for key, (a, b) in span.items():
        v = ships.get(imo.get(key))
        if v and v["teu_nominal"] >= 10000 and fy0 <= a < fy1:
            calls.append((key, a, b, v))
    # sailings
    sails = defaultdict(list)
    for r in read(f("pilot_station_sailings_fy2026.csv")):
        name = re.sub(r"\s+", " ", re.sub(r"^M/V ", "", r["VESSEL"].strip().upper()))
        d = max(int(x.split("-")[0]) * 12 + int(x.split("-")[1]) for x in (r["DRAFT_FWD"], r["DRAFT_AFT"]))
        sails[BERTHS[r["FROM_BERTH"].rsplit(" ", 1)[0]]].append((eastern(r["SAILED"]), eastern(r["NOTICE_RECEIVED"]),
                                                                 name, r["CALL_SIGN"], d))
    # windows
    hw = []
    for line in open(f("salt_point_tide_predictions_2025_2026.txt"), encoding="utf-8"):
        p = line.rstrip("\n").split("\t")
        if len(p) == 6 and p[5] == "H":
            hw.append(dt.datetime.strptime(p[0] + " " + p[2], "%Y/%m/%d %I:%M %p").replace(tzinfo=STD).timestamp())
    hw.sort()
    passage = {"SLT": 110 * 60, "MRW": 160 * 60}
    tide = {t: [(h - 5400 - passage[t], h + 1800 - passage[t]) for h in hw] for t in passage}
    days = []
    d = dt.date(2025, 6, 1)
    while d < dt.date(2026, 8, 1):
        days.append((dt.datetime(d.year, d.month, d.day, 6, tzinfo=ET).timestamp(),
                     dt.datetime(d.year, d.month, d.day, 20, tzinfo=ET).timestamp()))
        d += dt.timedelta(days=1)
    saved = {t: 0.0 for t in TERMS}
    per = {}
    for (term, visit), a, b, v in calls:
        names = {v["name"].upper()} | {x["name"].upper() for x in v["former_names"]}
        s = min((x for x in sails[term] if b <= x[0] <= b + 40 * 3600 and (x[2] in names or x[3] == v["call_sign"])),
                key=lambda x: x[0])
        sailed, notice, _, _, draft = s
        cut = (b - a) / 5
        if term in passage and draft >= 41 * 12:
            sv = sailed - first_open(notice + 3600 - cut, tide[term])
        elif term not in passage and v["loa_m"] * 3.28084 >= 1250:
            sv = sailed - first_open(notice + 3600 - cut, days)
        else:
            sv = cut
        saved[term] += sv / 3600
        per[visit] = sv / 3600

    def place(crane, t):
        last = where[crane][0][1]
        for u, term in where[crane]:
            if u > t:
                break
            last = term
        return last
    # ask A: each crane's time out on any breakdown order or SRT breakdown job, overlaps once
    types = {r["WONUM"]: r for r in read(f("cmms_sts_work_orders_fy2026.csv"))}
    down = {}
    stops = defaultdict(list)
    for r in read(f("cmms_asset_status_history_fy2026.csv")):
        t = eastern(r["CHANGEDATE"])
        k = (r["ASSETNUM"], r["WONUM"])
        if r["ISRUNNING"] == "0":
            down[k] = t
        else:
            a = down.pop(k)
            if r["WONUM"] and types[r["WONUM"]]["WORKTYPE"] == "BD":
                stops[by_asset[r["ASSETNUM"]]].append((a, t))
    jobs = sheet(f("srt_maintenance_jobs_2025-26.xlsx"), "Jobs")
    for j in jobs:
        if j["JOB_CLASS"] == "BRKDN":
            stops[by_tag[j["EQUIP_TAG"]]].append((j["DOWN_UTC"].replace(tzinfo=UTC).timestamp(),
                                                  j["UP_UTC"].replace(tzinfo=UTC).timestamp()))
    ask_a = {t: 0.0 for t in TERMS}
    for crane, ps in stops.items():
        merged = []
        for a, b in sorted(ps):
            if merged and a <= merged[-1][1]:
                merged[-1][1] = max(merged[-1][1], b)
            else:
                merged.append([a, b])
        for a, b in merged:
            ask_a[place(crane, a)] += max(0.0, min(b, fy1) - max(a, fy0)) / 3600
    # ask B
    noon = lambda d: dt.datetime(d.year, d.month, d.day, 12, tzinfo=ET).timestamp()
    d0, d1 = dt.date(2025, 7, 1), dt.date(2026, 6, 30)
    ask_b = {t: 0.0 for t in TERMS}
    contract = defaultdict(list)
    for r in read(f("cmms_labor_transactions_fy2026.csv")):
        o = types.get(r["WONUM"])
        if o is None:
            continue
        d = dt.date.fromisoformat(r["TRANSDATE"])
        if r["LABORCODE"] == "CONTRACT":
            contract[(r["WONUM"], d)].append(float(r["REGULARHRS"]))
        if d0 <= d <= d1:
            ask_b[place(by_asset[o["ASSETNUM"]], noon(d))] += float(r["REGULARHRS"])
    spans = {(by_asset[o["ASSETNUM"]], eastern(o["OUT_OF_SERVICE"]), eastern(o["RETURNED_TO_SERVICE"]))
             for o in types.values() if o["OUT_OF_SERVICE"]}
    jobcrane = {}
    for j in jobs:
        crane = by_tag[j["EQUIP_TAG"]]
        jobcrane[j["JOB_NO"]] = crane
        if j["DOWN_UTC"] is not None and (crane, j["DOWN_UTC"].replace(tzinfo=UTC).timestamp(),
                                          j["UP_UTC"].replace(tzinfo=UTC).timestamp()) in spans:
            continue
        t = (j["DOWN_UTC"] or j["RAISED_UTC"]).replace(tzinfo=UTC).timestamp()
        d = dt.datetime.fromtimestamp(t, ET).date()
        if d0 <= d <= d1:
            ask_b[place(crane, t)] += float(j["LABOUR_HRS"])
    sheet_rows = sheet(f("crane_maker_service_timesheets_fy2026.xlsx"), "Timesheets")
    # a CONTRACT line keyed from a service report of several days carries the report's hours on its first day
    report = defaultdict(list)
    for s in sheet_rows:
        m = re.search(r"\d{7}", str(s["CUSTOMER_REF"]))
        if m and m.group(0) in types:
            report[(m.group(0), s["SERVICE_REPORT"])].append((dt.date.fromisoformat(str(s["WORK_DATE"])[:10]),
                                                             float(s["HOURS"])))
    whole = set()
    for (wn, rep), days in report.items():
        if len(days) > 1:
            first, total = min(days)[0], sum(h for _, h in days)
            if total in contract[(wn, first)]:
                contract[(wn, first)].remove(total)
                whole.add((wn, rep))
    for s in sheet_rows:
        d = dt.date.fromisoformat(str(s["WORK_DATE"])[:10])
        if not (d0 <= d <= d1):
            continue
        h = float(s["HOURS"])
        m = re.search(r"\d{7}", str(s["CUSTOMER_REF"]))
        if m and m.group(0) in types:
            if (m.group(0), s["SERVICE_REPORT"]) in whole:
                continue
            if h in contract[(m.group(0), d)]:
                contract[(m.group(0), d)].remove(h)
                continue
            crane = by_asset[types[m.group(0)]["ASSETNUM"]]
        else:
            crane = jobcrane[str(s["CUSTOMER_REF"]).strip()]
        ask_b[place(crane, noon(d))] += h
    # compare with the goldens
    wb = openpyxl.load_workbook(os.path.join(golden, "ternmouth_2027_crane_placement.xlsx"), read_only=True)
    bh = {r[0]: r[3] for r in list(wb["Berth hours saved"].iter_rows(values_only=True))[1:7]}
    cc = {r[0]: (r[1], r[2]) for r in list(wb["Crane condition"].iter_rows(values_only=True))[1:7]}
    rec = {r[0]: r[1] for r in wb["Recommendation"].iter_rows(values_only=True) if r and r[0]}
    wb.close()
    rank = sorted(TERMS, key=lambda t: -saved[t])
    bad = []
    for t in TERMS:
        if bh[TERMS[t]] != round(saved[t]):
            bad.append(("saved", t, bh[TERMS[t]], saved[t]))
        if cc[TERMS[t]] != (round(ask_a[t]), round(ask_b[t])):
            bad.append(("asks", t, cc[TERMS[t]], round(ask_a[t], 2), round(ask_b[t], 2)))
    if rec["Recommended terminal"] != TERMS[rank[0]] or rec["Runner-up"] != TERMS[rank[1]]:
        bad.append(("call", rec["Recommended terminal"], rec["Runner-up"], rank[:2]))
    if rec["Lead over the runner-up, berth hours"] != round(saved[rank[0]]) - round(saved[rank[1]]):
        bad.append(("lead", rec["Lead over the runner-up, berth hours"]))
    with open(os.path.join(golden, "ternmouth_fy2026_large_vessel_calls.csv"), encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    if len(rows) != len(calls) or any(abs(float(r["BERTH_HOURS_SAVED"]) - per[r["VISIT"]]) > 0.00501 for r in rows):
        bad.append(("calls table",))
    print("verify: %s first (%d), %s second (%d); %d calls; %s" % (
        TERMS[rank[0]], round(saved[rank[0]]), TERMS[rank[1]], round(saved[rank[1]]), len(calls),
        "goldens match" if not bad else "MISMATCH %s" % bad))
    if bad:
        raise SystemExit(1)
    return dict(saved=saved, ask_a=ask_a, ask_b=ask_b)


if __name__ == "__main__":
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "target"),
         sys.argv[2] if len(sys.argv) > 2 else os.path.join(here, "golden"))
