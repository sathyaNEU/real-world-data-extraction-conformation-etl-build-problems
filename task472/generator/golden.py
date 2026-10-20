"""The golden deliverables, built from the shipped pack only (bundle.py): the placement workbook with each terminal's
berth hours saved, the chart and the two crane-condition tables, the conformed fiscal year 2026 large-vessel call
table, and generator/figures.json."""
import csv
import datetime as dt
import json
import os
from collections import defaultdict
from zoneinfo import ZoneInfo

import bundle as B

ET = ZoneInfo("America/New_York")
WB = "ternmouth_2027_crane_placement.xlsx"
CALLS = "ternmouth_fy2026_large_vessel_calls.csv"


def et(t):
    return dt.datetime.fromtimestamp(t, ET).strftime("%Y-%m-%d %H:%M")


def compute(target):
    rows = B.call_table(target, "gold")
    saved = {t: 0.0 for t in B.TERMS}
    ncalls, held, cut = defaultdict(int), defaultdict(int), defaultdict(float)
    for r in rows:
        saved[r["term"]] += r["saved"]
        ncalls[r["term"]] += 1
        held[r["term"]] += r["gate"] is not None and r["sailed"] > r["notice"] + B.NOTICE
        cut[r["term"]] += r["cut"]
    rank = B.ranking(saved)
    return dict(rows=rows, saved=saved, ncalls=dict(ncalls), held=dict(held), cut=dict(cut), rank=rank,
                a=B.breakdown_hours(target), b=B.labor_hours(target))


def figures(g):
    n = B.NAME
    top, second = g["rank"][0], g["rank"][1]
    whole = {t: int(round(g["saved"][t])) for t in B.TERMS}
    return {
        "call": {"terminal": n[top]},
        "components": {
            "berth_hours_saved": {n[t]: whole[t] for t in B.TERMS},
            "runner_up": n[second],
            "gap_hours": whole[top] - whole[second],
            "large_vessel_calls": {n[t]: g["ncalls"][t] for t in B.TERMS},
            "total_large_vessel_calls": sum(g["ncalls"].values()),
        },
        "a_breakdown_hours": {n[t]: int(round(g["a"][t])) for t in B.TERMS},
        "b_labor_hours": {n[t]: int(round(g["b"][t])) for t in B.TERMS},
        "inputs": {"teu_line": B.TEU_LINE, "fiscal_year": "1 July 2025 to 30 June 2026", "deep_draft_ft": 41,
                   "window_before": "1 h 30 min", "window_after": "30 min", "long_ft": 1250,
                   "daylight": ["0600", "2000"], "passage_minutes": B.PASSAGE_MIN,
                   "rules": ["3.2", "3.3", "3.6", "8.1", "9.2", "9.3", "3.7"]},
    }


def write_calls_csv(g, path):
    cols = ["TERMINAL", "VISIT", "SHIP", "IMO_NUMBER", "FIRST_CONTAINER_MOVE", "LAST_CONTAINER_MOVE",
            "BERTH_HOURS_SAVED"]
    out = []
    for r in g["rows"]:
        out.append([B.NAME[r["term"]], r["visit"], r["vessel"]["name"], r["imo"], et(r["first"]), et(r["last"]),
                    "%.2f" % r["saved"]])
    out.sort(key=lambda z: (z[4], z[0]))
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(cols)
        w.writerows(out)
    return len(out)


def write_workbook(g, f, path):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.marker import DataPoint
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    n = B.NAME
    wb = Workbook()
    ws = wb.active
    ws.title = "Recommendation"
    bold = Font(name="Calibri", bold=True)
    head = PatternFill("solid", fgColor="DCE6F0")
    rule = Border(bottom=Side(style="thin", color="8EA2B8"))
    top, second = g["rank"][0], g["rank"][1]
    bs = f["components"]["berth_hours_saved"]
    gap = f["components"]["gap_hours"]
    ws["A1"] = "Place the 2027 cranes at %s: %d berth hours saved on FY2026's large-vessel calls" % (n[top], bs[n[top]])
    ws["A1"].font = Font(name="Calibri", bold=True, size=13, color="1F3B5C")
    block = [("To", "Board of Commissioners"),
             ("From", "Jon Hernandez, Director, Terminal Equipment"),
             ("Date", "14 October 2026, for the meeting of 27 October 2026"),
             ("Subject", "Crane Replacement Program, 2027 order (six ship-to-shore cranes): receiving terminal"),
             ("Reference", "CRP-2027-01")]
    r = 3
    for a, b in block:
        ws.cell(row=r, column=1, value=a).font = bold
        ws.cell(row=r, column=2, value=b)
        r += 1
    for c in range(1, 3):
        ws.cell(row=r - 1, column=c).border = rule
    r += 1
    rows = [("Recommended terminal", n[top]),
            ("Berth hours saved, FY2026 large-vessel calls", bs[n[top]]),
            ("Runner-up", n[second]),
            ("Runner-up berth hours saved", bs[n[second]]),
            ("Lead over the runner-up, berth hours", gap)]
    for a, b in rows:
        ws.cell(row=r, column=1, value=a).font = bold
        cell = ws.cell(row=r, column=2, value=b)
        if isinstance(b, int):
            cell.number_format = "#,##0"
            cell.alignment = Alignment(horizontal="left")
        r += 1
    ws.cell(row=r - 5, column=2).font = Font(name="Calibri", bold=True, size=12)
    r += 1
    prose = (
        "Rule 3.6 counts the hours by which each call, all fast to let go, would have been shorter had its cargo work "
        "taken four fifths of the time. A ship lets go at the sailing time the pilot station gives it: the first time an "
        "hour or more after its ready-to-sail notice that the Harbor Master's Directions allow. Ships drawing 41 feet or "
        "more leave the Saltworks and Millrace berths only in the window around high water at Salt Point (9.2, 9.3), and "
        "ships of 1,250 feet or more leave the harbor berths only from 0600 to 2000 (8.1). So each call's saving is its "
        "sailing time less the sailing time it would have been given from a notice a fifth of its cargo work earlier. "
        "Saltworks' deep ships mostly finish ahead of their tide, so a faster finish would only have lengthened their "
        "wait; its year comes to %d hours. Millrace's deep ships mostly finish just after the tide they were worked to "
        "has gone and wait for the next one; four fifths of the cargo time would have put them on the earlier tide: %d "
        "hours, %d ahead of %s." % (bs["Saltworks"], bs[n[top]], gap, n[second]))
    ws.cell(row=r, column=1, value="Finding").font = bold
    ws.cell(row=r, column=2, value=prose).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 160
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 96

    ws2 = wb.create_sheet("Berth hours saved")
    hdr = ["Terminal", "FY2026 large-vessel calls", "Sailings held for the tide or daylight", "Berth hours saved"]
    ws2.append(hdr)
    order = sorted(B.TERMS, key=lambda t: -g["saved"][t])
    for t in order:
        ws2.append([n[t], g["ncalls"][t], g["held"][t], bs[n[t]]])
    ws2.append(["All terminals", sum(g["ncalls"].values()), sum(g["held"].values()),
                int(round(sum(g["saved"].values())))])
    for c in ws2[1]:
        c.font = bold
        c.fill = head
        c.alignment = Alignment(wrap_text=True, vertical="center")
    last = len(order) + 2
    for c in ws2[last]:
        c.font = bold
        c.border = Border(top=Side(style="thin", color="1B1F24"))
    for row in ws2.iter_rows(min_row=2, max_row=last, min_col=2, max_col=4):
        for c in row:
            c.number_format = "#,##0"
    ws2.row_dimensions[1].height = 32
    for col, wd in zip("ABCD", (18, 15, 19, 14)):
        ws2.column_dimensions[col].width = wd
    ws2.freeze_panes = "A2"
    ws2.cell(row=last + 2, column=1, value="Whole hours; each terminal's calls added unrounded (Rule 3.7), and the "
                                           "total from the unrounded terminal figures. Source: Authority and SRT TOS "
                                           "crane hours, pilot station sailings, Salt Point tide predictions, vessel "
                                           "register, Harbor Master's Directions 2025.")
    ws2.cell(row=last + 2, column=1).font = Font(name="Calibri", italic=True, size=9, color="555555")
    ws2.cell(row=last + 2, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws2.merge_cells(start_row=last + 2, start_column=1, end_row=last + 2, end_column=4)
    ws2.row_dimensions[last + 2].height = 52
    ch = BarChart()
    ch.type = "bar"
    ch.title = "%s's large-vessel calls would have saved the most berth time, FY2026" % n[top]
    ch.y_axis.title = "Berth hours saved"
    ch.x_axis.title = None
    data = Reference(ws2, min_col=4, min_row=1, max_row=len(order) + 1)
    cats = Reference(ws2, min_col=1, min_row=2, max_row=len(order) + 1)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ch.legend = None
    ser = ch.series[0]
    ser.graphicalProperties.solidFill = "A9AFB7"
    ser.graphicalProperties.line.noFill = True
    hi = DataPoint(idx=order.index(top))
    hi.graphicalProperties.solidFill = "2A78D6"
    hi.graphicalProperties.line.noFill = True
    ser.dPt = [hi]
    ser.dLbls = DataLabelList()
    ser.dLbls.showVal = True
    ser.dLbls.showSerName = False
    ser.dLbls.showCatName = False
    ser.dLbls.showLegendKey = False
    ch.gapWidth = 60
    ch.height, ch.width = 9.6, 16.5
    ch.x_axis.tickLblSkip = 1
    ch.x_axis.scaling.orientation = "maxMin"
    ch.y_axis.crosses = "max"
    ch.y_axis.majorGridlines = None
    ch.y_axis.number_format = "#,##0"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ws2.add_chart(ch, "F2")

    ws3 = wb.create_sheet("Crane condition")
    ws3.append(["Terminal", "Hours STS cranes stood out of service for breakdowns, FY2026",
                "Maintenance labor hours on STS cranes, FY2026 (staff and contractors)"])
    for t in B.TERMS:
        ws3.append([n[t], f["a_breakdown_hours"][n[t]], f["b_labor_hours"][n[t]]])
    ws3.append(["All terminals", int(round(sum(g["a"].values()))), int(round(sum(g["b"].values())))])
    for c in ws3[1]:
        c.font = bold
        c.fill = head
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for c in ws3[8]:
        c.font = bold
        c.border = Border(top=Side(style="thin", color="1B1F24"))
    ws3.row_dimensions[1].height = 46
    for col, wd in zip("ABC", (18, 26, 28)):
        ws3.column_dimensions[col].width = wd
    for row in ws3.iter_rows(min_row=2, max_row=8, min_col=2, max_col=3):
        for c in row:
            c.number_format = "#,##0"
    ws3.freeze_panes = "A2"

    ws4 = wb.create_sheet("Notes")
    notes = [
        ("Large vessel", "Nominal capacity of 10,000 TEU or more in the vessel register (Rule 3.2)."),
        ("Call and year", "A ship's stay alongside one terminal, all fast to let go (Rule 3.1); in FY2026 when its "
                          "first container move fell from 1 July 2025 to 30 June 2026, Eastern time (Rule 3.3). "
                          "SRT records UTC and was converted."),
        ("Sailing", "Each call's sailing from the pilot station's record (its ready-to-sail notice and sailing time), "
                    "matched by ship name, former name or call sign at the call's berth."),
        ("Berth hours saved", "Per call: the sailing time less the first time, an hour or more after the notice brought "
                              "forward by a fifth of the cargo work, at which the Directions allow the ship to sail: "
                              "Salt Point within 1 h 30 min before to 30 min after high water for drafts of 41 feet "
                              "or more (9.2), on the passage times in 9.3, and 0600 to 2000 for ships of 1,250 feet or "
                              "more in the harbor (8.1). Tide predictions are standard time; an hour added in daylight "
                              "saving time."),
        ("Out of service", "Breakdown periods (order type BD at the pull) from the asset status history and the "
                           "SRT log; a reopened order's periods each; a crane held out on two breakdown orders at once "
                           "counted once, as are February 2026's river stops held in both systems; each stop at the "
                           "terminal where its crane was working at the time (three cranes moved along the common rail "
                           "during the year)."),
        ("Labor", "Labor ledger hours on crane orders, SRT job hours not also in the Authority's system, and the "
                  "makers' timesheet days not already keyed into the ledger (one CONTRACT line per service report, "
                  "dated the visit's first day); each at the terminal where the crane was working that day; yard "
                  "equipment left out."),
        ("Let go", "A call ends at the sailing time in the pilot station's record. The AIS berth stretches close a "
                   "few minutes later, when the ship's reports leave the berth polygon, and were not used."),
        ("Clock change", "On 2 November 2025 the Authority TOS writes the 01:00 hour twice; its rows were read in "
                         "file order, the first occurrence first."),
        ("As of", "Extracts pulled 2 to 6 October 2026 (export manifest)."),
    ]
    ws4.append(["Item", "Basis"])
    for a, b in notes:
        ws4.append([a, b])
    for c in ws4[1]:
        c.font = bold
        c.fill = head
    for row in ws4.iter_rows(min_row=2, max_row=len(notes) + 1, min_col=2, max_col=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws4.iter_rows(min_row=2, max_row=len(notes) + 1, min_col=1, max_col=1):
        for c in row:
            c.font = bold
            c.alignment = Alignment(vertical="top")
    ws4.column_dimensions["A"].width = 18
    ws4.column_dimensions["B"].width = 100
    for i in range(2, len(notes) + 2):
        ws4.row_dimensions[i].height = 48
    for sh in (ws, ws2, ws3, ws4):
        sh.page_setup.orientation = "landscape"
        sh.page_setup.paperSize = sh.PAPERSIZE_LETTER
        sh.sheet_properties.pageSetUpPr.fitToPage = True
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
    ws2.print_title_rows = "1:1"
    wb.save(path)


def main(target, gdir, fig_path):
    os.makedirs(gdir, exist_ok=True)
    g = compute(target)
    f = figures(g)
    write_workbook(g, f, os.path.join(gdir, WB))
    import postprocess as PP
    PP.normalize_ooxml(os.path.join(gdir, WB), "Jon Hernandez", "Tanya Abbott", "2026-10-07T13:05:00Z",
                       "2026-10-14T20:41:16Z", excel=True)
    write_calls_csv(g, os.path.join(gdir, CALLS))
    with open(fig_path, "w", encoding="utf-8") as fh:
        json.dump(f, fh, indent=1, sort_keys=False)
        fh.write("\n")
    return g, f
