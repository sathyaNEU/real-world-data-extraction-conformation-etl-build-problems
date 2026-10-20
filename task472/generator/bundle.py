"""Read the shipped pack and build the program's figures from it: the fiscal year 2026 large-vessel calls from both
terminal operating systems and the vessel register, each call's sailing from the pilot station's record, the Salt
Point windows from the tide predictions and the Harbor Master's directions, each call's berth hours saved, and the two
asks (hours cranes stood out of service for breakdowns, maintenance labor hours). Only the pack's files are read. The
options replay the natural mishandlings and the wrong routes; checks.py and ceiling.py run them."""
import csv
import datetime as dt
import json
import os
import re
from collections import defaultdict
from zoneinfo import ZoneInfo

import numpy as np
import openpyxl

ET = ZoneInfo("America/New_York")
UTC = dt.timezone.utc
LST = dt.timezone(dt.timedelta(hours=-5))
H = 3600.0
FILES = dict(tos="tos_crane_activity_2026.csv", visits="tos_vessel_visits_2026.csv",
             rvt="srt_tos_export_fy2026.xlsx", pilot="pilot_station_sailings_fy2026.csv",
             tides="salt_point_tide_predictions_2025_2026.txt", vessels="vessel_register_extract_2026-10-05.json",
             register="sts_crane_register.csv", wo="cmms_sts_work_orders_fy2026.csv",
             status="cmms_asset_status_history_fy2026.csv", labor="cmms_labor_transactions_fy2026.csv",
             oem="crane_maker_service_timesheets_fy2026.xlsx", jobs="srt_maintenance_jobs_2025-26.xlsx")
FY0 = dt.datetime(2025, 7, 1, tzinfo=ET).timestamp()
FY1 = dt.datetime(2026, 7, 1, tzinfo=ET).timestamp()
D0, D1 = dt.date(2025, 7, 1), dt.date(2026, 6, 30)
TERMS = ["ORD", "PNC", "GNP", "HLY", "SLT", "MRW"]
NAME = dict(ORD="Ore Dock", PNC="Pine Neck", GNP="Gannet Point", HLY="Halyard", SLT="Saltworks", MRW="Millrace")
BERTH_TERM = {"ORE DOCK": "ORD", "PINE NECK": "PNC", "GANNET PT": "GNP", "HALYARD": "HLY", "SALTWORKS": "SLT",
              "MILLRACE": "MRW"}
# the directions: 8.1 (1,250 feet, 0600 to 2000 local), 9.2 (41 feet, Salt Point from high water less 1 h 30 to
# plus 0 h 30), 9.3 (passage times), 3.4 (an hour's notice); the program: 3.2 (10,000 TEU), 3.5 (four fifths)
TEU_LINE = 10000
LONG_M = 1250 * 0.3048
DEEP_IN = 41 * 12
WIN = (-90 * 60, 30 * 60)
PASSAGE = dict(SLT=110 * 60, MRW=160 * 60)
PASSAGE_MIN = dict(Saltworks=110, Millrace=160)
DAYLIGHT = (6, 20)
NOTICE = 3600
CUT = 0.2
CREW = {"1": "ORD", "2": "PNC", "3": "GNP", "4": "HLY", "6": "SLT", "7": "MRW"}   # the technicians' numbering

_CACHE = {}


def p(target, key):
    return os.path.join(target, FILES[key])


def rows_csv(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def et_ts(s, fold=0):
    d = dt.datetime.strptime(s, "%Y-%m-%d %H:%M:%S" if len(s) > 16 else "%Y-%m-%d %H:%M")
    return d.replace(tzinfo=ET, fold=fold).timestamp()


def utc_ts(d):
    return d.replace(tzinfo=UTC).timestamp()


def cached(key, fn):
    if key not in _CACHE:
        _CACHE[key] = fn()
    return _CACHE[key]


# ---------------------------------------------------------------------------------------------- readers
def vessels(target):
    vj = json.load(open(p(target, "vessels"), encoding="utf-8"))
    return {int(v["imo"]): v for v in vj["vessels"]}


def register(target):
    return {r["CRANE_ID"]: r for r in rows_csv(p(target, "register"))}


def crane_hours(target):
    """Every crane-hour in both terminal systems: {(term, visit): {crane: dict(first, last)}}, each crane's working
    hours as (UTC hour, terminal), and each visit's IMO number."""
    def build():
        runs = defaultdict(dict)
        where = defaultdict(list)
        seen = defaultdict(int)
        for r in rows_csv(p(target, "tos")):
            key = (r["TERMINAL"], r["VISIT_ID"])
            lab = (key, r["CRANE_ID"], r["HOUR_START"])
            fold = 1 if seen[lab] else 0
            seen[lab] += 1
            f, last = et_ts(r["FIRST_MOVE"], fold), et_ts(r["LAST_MOVE"], fold)
            d = runs[key].setdefault(r["CRANE_ID"], dict(first=f, last=last))
            d["first"], d["last"] = min(d["first"], f), max(d["last"], last)
            where[r["CRANE_ID"]].append((f, r["TERMINAL"]))
        imo = {}
        for r in rows_csv(p(target, "visits")):
            imo[(r["TERMINAL"], r["VISIT_ID"])] = int(r["IMO_NO"])
        tag = {}
        if os.path.exists(p(target, "register")):
            tag = {r["SRT_TAG"]: k for k, r in register(target).items() if r["SRT_TAG"]}
        wb = openpyxl.load_workbook(p(target, "rvt"), read_only=True)
        calls = list(wb["Calls"].iter_rows(values_only=True))
        for row in calls[1:]:
            rr = dict(zip(calls[0], row))
            imo[(rr["CALL_REF"][:3], rr["CALL_REF"])] = int(rr["LLOYDS"])
        qc = list(wb["QC Hours"].iter_rows(values_only=True))
        for row in qc[1:]:
            rr = dict(zip(qc[0], row))
            key = (rr["CALL_REF"][:3], rr["CALL_REF"])
            f, last = utc_ts(rr["FIRST_LIFT_UTC"]), utc_ts(rr["LAST_LIFT_UTC"])
            cr = tag.get(rr["QC"], rr["QC"])
            d = runs[key].setdefault(cr, dict(first=f, last=last))
            d["first"], d["last"] = min(d["first"], f), max(d["last"], last)
            where[cr].append((f, key[0]))
        wb.close()
        for k in where:
            where[k].sort()
        return runs, imo, where
    return cached(("hours", os.path.abspath(target)), build)


def large_calls(target, teu_line=TEU_LINE):
    """The fiscal year 2026 large-vessel calls: [dict(term, visit, imo, first, last, vessel)]."""
    runs, imo, _ = crane_hours(target)
    ves = vessels(target)
    out = []
    for key, cr in sorted(runs.items()):
        v = ves.get(imo.get(key))
        if v is None or v["teu_nominal"] < teu_line:
            continue
        first = min(d["first"] for d in cr.values())
        last = max(d["last"] for d in cr.values())
        if FY0 <= first < FY1:
            out.append(dict(term=key[0], visit=key[1], imo=v["imo"], first=first, last=last, vessel=v))
    return out


def norm(name):
    s = re.sub(r"^M/V\s+", "", name.strip().upper())
    return re.sub(r"\s+", " ", s)


def sailings(target):
    def build():
        out = []
        for r in rows_csv(p(target, "pilot")):
            berth = r["FROM_BERTH"].rsplit(" ", 1)[0]
            dr = []
            for col in ("DRAFT_FWD", "DRAFT_AFT"):
                f, i = r[col].split("-")
                dr.append(int(f) * 12 + int(i))
            out.append(dict(job=r["JOB_NO"], term=BERTH_TERM[berth], name=norm(r["VESSEL"]), cs=r["CALL_SIGN"].strip(),
                            notice=et_ts(r["NOTICE_RECEIVED"]), sailed=et_ts(r["SAILED"]), draft_in=max(dr)))
        return out
    return cached(("sail", os.path.abspath(target)), build)


def link(target):
    """Each large-vessel call's sailing: the first sailing from the call's terminal after its last move by the ship
    (its name at the time, current or former, or its call sign)."""
    def build():
        by_term = defaultdict(list)
        for s in sailings(target):
            by_term[s["term"]].append(s)
        out = {}
        for c in large_calls(target):
            v = c["vessel"]
            names = {norm(v["name"])} | {norm(f["name"]) for f in v.get("former_names", [])}
            cand = [s for s in by_term[c["term"]] if c["last"] <= s["sailed"] <= c["last"] + 40 * H
                    and (s["name"] in names or s["cs"] == v["call_sign"])]
            cand.sort(key=lambda s: s["sailed"])
            out[c["visit"]] = cand[0] if cand else None
        return out
    return cached(("link", os.path.abspath(target)), build)


def high_waters(target, as_clock=False):
    """Predicted high waters at Salt Point as UTC seconds; as_clock reads the standard-time table as clock time."""
    out = []
    for line in open(p(target, "tides"), encoding="utf-8"):
        f = line.rstrip("\n").split("\t")
        if len(f) == 6 and f[5] == "H":
            d = dt.datetime.strptime(f[0] + " " + f[2], "%Y/%m/%d %I:%M %p")
            out.append(d.replace(tzinfo=ET if as_clock else LST).timestamp())
    return np.array(sorted(out))


def windows(target, term, passage=True, as_clock=False):
    hw = high_waters(target, as_clock)
    pp = PASSAGE[term] if passage else 0
    return np.stack([hw + WIN[0] - pp, hw + WIN[1] - pp], axis=1)


def day_windows():
    out = []
    d = dt.date(2025, 6, 1)
    while d <= dt.date(2026, 8, 1):
        out.append((dt.datetime(d.year, d.month, d.day, DAYLIGHT[0], tzinfo=ET).timestamp(),
                    dt.datetime(d.year, d.month, d.day, DAYLIGHT[1], tzinfo=ET).timestamp()))
        d += dt.timedelta(days=1)
    return np.array(out)


def earliest(t, wins):
    i = int(np.searchsorted(wins[:, 1], t))
    return float(max(t, wins[i, 0]))


def gate(c, s):
    """Which window binds a call's sailing: 'tide' for a river sailing at 41 feet or more, 'day' for a harbor ship of
    1,250 feet or more, None otherwise."""
    if c["term"] in ("SLT", "MRW"):
        return "tide" if s["draft_in"] >= DEEP_IN else None
    return "day" if c["vessel"]["loa_m"] >= LONG_M else None


# ---------------------------------------------------------------------------------------------- berth hours saved
def call_table(target, route="gold"):
    """One row per large-vessel call with its berth hours saved under a route.
    gold: the sailing re-allocated from the earlier readiness (notice less a fifth of the cargo work, plus the hour's
    notice) to the first time the directions then permit;
    naive: a fifth of the cargo work; tide_zero: held river sailings save nothing, the rest the cut;
    gated_zero: every held sailing (tide or daylight) saves nothing; waited_zero: a ship that waited past its
    readiness saves nothing; daylight_missed: the replay for the river only; no_passage: the window at the berth;
    lst_clock: the tide table read as clock time; ready_last: readiness at the last move less the cut."""
    key = ("table", os.path.abspath(target), route)
    if key in _CACHE:
        return [dict(r) for r in _CACHE[key]]
    lk = link(target)
    dayw = day_windows()
    tide = {t: windows(target, t) for t in ("SLT", "MRW")}
    rows = []
    for c in large_calls(target):
        s = lk[c["visit"]]
        cut = CUT * (c["last"] - c["first"])
        g = gate(c, s)
        ready = s["notice"] + NOTICE
        if route == "naive" or g is None:
            sv = cut
        elif route == "tide_zero":
            sv = 0.0 if g == "tide" else cut
        elif route == "gated_zero":
            sv = 0.0
        elif route == "waited_zero":
            sv = 0.0 if s["sailed"] > ready else cut
        elif route == "daylight_missed" and g == "day":
            sv = cut
        else:
            if g == "tide":
                if route == "no_passage":
                    W = windows(target, c["term"], passage=False)
                elif route == "lst_clock":
                    W = windows(target, c["term"], as_clock=True)
                else:
                    W = tide[c["term"]]
            else:
                W = dayw
            r_cf = (c["last"] - cut) if route == "ready_last" else ready - cut
            sv = s["sailed"] - earliest(r_cf, W)
        rows.append(dict(c, sailed=s["sailed"], notice=s["notice"], gate=g, cut=cut / H, saved=sv / H))
    _CACHE[key] = rows
    return [dict(r) for r in rows]


def totals(target, route="gold"):
    t = {x: 0.0 for x in TERMS}
    for r in call_table(target, route):
        t[r["term"]] += r["saved"]
    return t


def ranking(tot):
    return sorted(tot, key=lambda k: -tot[k])


# ---------------------------------------------------------------------------------------------- where cranes stood
def located(target, route="truth"):
    """A function (crane, UTC seconds) -> terminal. truth: the terminal of the crane's working hours on either side
    (the last before, else the first after); current: the crane register's terminal."""
    reg = register(target)
    if route == "current":
        return lambda crane, t: reg[crane]["TERMINAL"]
    _, _, where = crane_hours(target)
    idx = {k: (np.array([a for a, _ in v]), [b for _, b in v]) for k, v in where.items()}

    def f(crane, t):
        ts_, terms = idx[crane]
        i = int(np.searchsorted(ts_, t, side="right")) - 1
        return terms[max(i, 0)]
    return f


def crane_keys(target):
    reg = register(target)
    asset = {r["ASSETNUM"]: k for k, r in reg.items()}
    tag = {r["SRT_TAG"]: k for k, r in reg.items() if r["SRT_TAG"]}
    return asset, tag


def jobs(target):
    def build():
        wb = openpyxl.load_workbook(p(target, "jobs"), read_only=True)
        rows = list(wb["Jobs"].iter_rows(values_only=True))
        wb.close()
        return [dict(zip(rows[0], r)) for r in rows[1:]]
    return cached(("jobs", os.path.abspath(target)), build)


def overlap(a, b, s, e):
    return max(0.0, min(b, e) - max(a, s))


# ---------------------------------------------------------------------------------------------- ask A
def breakdown_hours(target, route="truth"):
    """Each terminal's hours of ship-to-shore crane breakdown out of service over fiscal year 2026.
    truth: the status history's out-of-service periods on orders whose type at the pull is BD, and SRT's BRKDN
    jobs, a stop held in both systems once, each at the terminal where its crane stood;
    current: placed at the crane register's terminal (the cracker's handling);
    both_counted: February's stops held in both systems counted twice;
    order_spans: the orders' OUT_OF_SERVICE to RETURNED_TO_SERVICE instead of the status history;
    code: the status history's BRKDWN code instead of the order type (damage kept)."""
    asset, tag = crane_keys(target)
    wtype = {r["WONUM"]: r["WORKTYPE"] for r in rows_csv(p(target, "wo"))}
    per = []
    if route == "order_spans":
        for r in rows_csv(p(target, "wo")):
            if r["WORKTYPE"] == "BD" and r["OUT_OF_SERVICE"]:
                per.append((asset[r["ASSETNUM"]], et_ts(r["OUT_OF_SERVICE"]), et_ts(r["RETURNED_TO_SERVICE"]), "a"))
    else:
        open_ = {}
        for r in rows_csv(p(target, "status")):
            t = et_ts(r["CHANGEDATE"])
            if r["ISRUNNING"] == "0":
                open_[r["ASSETNUM"]] = (t, r)
                continue
            a, r0 = open_.pop(r["ASSETNUM"])
            keep = (r0["CODE"] == "BRKDWN") if route == "code" else (wtype.get(r0["WONUM"]) == "BD")
            if keep:
                per.append((asset[r["ASSETNUM"]], a, t, "a"))
    for j in jobs(target):
        if j["JOB_CLASS"] == "BRKDN":
            per.append((tag[j["EQUIP_TAG"]], utc_ts(j["DOWN_UTC"]), utc_ts(j["UP_UTC"]), "r"))
    if route != "both_counted":
        keep, seen = [], set()
        for x in sorted(per, key=lambda z: (z[1], z[0], z[3])):
            k = (x[0], int(x[1]), int(x[2]))
            if k in seen:
                continue
            seen.add(k)
            keep.append(x)
        per = keep
    where = located(target, "current" if route == "current" else "truth")
    tot = {x: 0.0 for x in TERMS}
    for cr, a, b, src in per:
        tot[where(cr, a)] += overlap(a, b, FY0, FY1) / H
    return tot


# ---------------------------------------------------------------------------------------------- ask B
def oem_sheet(target):
    def build():
        wb = openpyxl.load_workbook(p(target, "oem"), read_only=True)
        rows = list(wb["Timesheets"].iter_rows(values_only=True))
        wb.close()
        return [dict(zip(rows[0], r)) for r in rows[1:]]
    return cached(("oem", os.path.abspath(target)), build)


def labor_hours(target, route="truth"):
    """Each terminal's maintenance labor hours on ship-to-shore cranes over fiscal year 2026 (work days in the year).
    truth: the labor ledger on crane orders, SRT's job hours for jobs not also in the Authority's system, and
    the makers' timesheet days not already keyed into the ledger (same order, day and hours), each at the terminal
    where the crane stood that day;
    all_timesheets: every timesheet day added to the ledger (the cracker's handling, with the register's terminal);
    ledger_only: no timesheet day added; both_counted: February's jobs in both systems counted twice;
    current: the crane register's terminal; with_yard: the ledger's yard equipment orders kept."""
    asset, tag = crane_keys(target)
    wo = {r["WONUM"]: r for r in rows_csv(p(target, "wo"))}
    where = located(target, "current" if route in ("current", "all_timesheets") else "truth")
    noon = lambda d: dt.datetime(d.year, d.month, d.day, 12, tzinfo=ET).timestamp()
    tot = {x: 0.0 for x in TERMS}
    keyed = defaultdict(list)
    for r in rows_csv(p(target, "labor")):
        d = dt.date.fromisoformat(r["TRANSDATE"])
        o = wo.get(r["WONUM"])
        if o is None:
            if route == "with_yard" and D0 <= d <= D1 and r["LABORCODE"][:1] == "4":
                tot[CREW[r["LABORCODE"][1]]] += float(r["REGULARHRS"])
            continue
        h = float(r["REGULARHRS"])
        if r["LABORCODE"] == "CONTRACT":
            keyed[(r["WONUM"], d)].append(h)
        if D0 <= d <= D1:
            tot[where(asset[o["ASSETNUM"]], noon(d))] += h
    # SRT's jobs, less those also in the Authority's system (the same crane out of service at the same times)
    dup = set()
    if route != "both_counted":
        spans = set()
        for r in rows_csv(p(target, "wo")):
            if r["OUT_OF_SERVICE"]:
                spans.add((asset[r["ASSETNUM"]], et_ts(r["OUT_OF_SERVICE"]), et_ts(r["RETURNED_TO_SERVICE"])))
        for j in jobs(target):
            if j["DOWN_UTC"] is not None:
                k = (tag[j["EQUIP_TAG"]], utc_ts(j["DOWN_UTC"]), utc_ts(j["UP_UTC"]))
                if k in spans:
                    dup.add(j["JOB_NO"])
    for j in jobs(target):
        if j["JOB_NO"] in dup:
            continue
        t = utc_ts(j["DOWN_UTC"] if j["DOWN_UTC"] is not None else j["RAISED_UTC"])
        d = dt.datetime.fromtimestamp(t, ET).date()
        if D0 <= d <= D1:
            tot[where(tag[j["EQUIP_TAG"]], t)] += float(j["LABOUR_HRS"])
    if route == "ledger_only":
        return tot
    jobcrane = {j["JOB_NO"]: tag[j["EQUIP_TAG"]] for j in jobs(target)}
    for s in oem_sheet(target):
        d = dt.date.fromisoformat(str(s["WORK_DATE"])[:10])
        if not (D0 <= d <= D1):
            continue
        h = float(s["HOURS"])
        ref = str(s["CUSTOMER_REF"])
        m = re.search(r"(\d{7})", ref)
        if m and m.group(1) in wo:
            wn = m.group(1)
            if route != "all_timesheets" and h in keyed.get((wn, d), []):
                keyed[(wn, d)].remove(h)
                continue
            cr = asset[wo[wn]["ASSETNUM"]]
        else:
            cr = jobcrane[ref.strip()]
        tot[where(cr, noon(d))] += h
    return tot
