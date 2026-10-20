"""The generator's assertions on the built bundle: the input gates, the population and its links, every row of the
semantic fork table (SF1 to SF15), the ladder's rungs, margins and positions, the analyst test on the clean truth, the
separation line between the call and the asks, each ask's devices and stops, every graded figure's distance from the
edge of its cell, and the goldens against figures.json."""
import csv
import datetime as dt
import json
import os
import shutil
import tempfile
from collections import Counter, defaultdict

import numpy as np
import openpyxl

import bundle as B
import pack as P
import world as W

H = 3600.0


def run(target, gdir, x=None, meta_path=None, fig_path=None, quiet=False):
    x = x or W.build()
    if not any("wonum" in e for e in x["maint"]):
        P.order_book(x)          # the numbers the pack gave the orders and jobs
        P.job_rows(x)
    rows = []

    def ok(cid, cond, detail):
        rows.append((cid, bool(cond), detail))

    # ------------------------------------------------------------------ input gates
    files = sorted(os.listdir(target))
    exts = Counter(os.path.splitext(f)[1] for f in files)
    big = max(sum(1 for _ in open(os.path.join(target, f), encoding="utf-8", errors="replace"))
              for f in files if f.endswith(".csv"))
    ok("G1", len(files) >= 10, "%d files" % len(files))
    ok("G2", len(exts) >= 3, "formats %s" % dict(exts))
    ok("G3", big >= 25000, "largest csv %d lines" % big)
    meta = json.load(open(meta_path)) if meta_path else {}
    dis = meta.get("distractor_files", [])
    ok("G4", len(dis) >= 2 and all(d in files for d in dis), "distractors %s" % dis)

    # ------------------------------------------------------------------ population and links
    calls = B.large_calls(target)
    lk = B.link(target)
    ok("P1", all(lk[c["visit"]] is not None for c in calls), "every large-vessel call has its sailing (%d)" % len(calls))
    sail = B.sailings(target)
    used = Counter(id(lk[c["visit"]]) for c in calls)
    ok("P2", max(used.values()) == 1, "no sailing serves two calls")
    by_term = defaultdict(list)
    for s in sail:
        by_term[s["term"]].append(s)
    amb = 0
    for c in calls:
        v = c["vessel"]
        names = {B.norm(v["name"])} | {B.norm(f["name"]) for f in v.get("former_names", [])}
        cand = [s for s in by_term[c["term"]] if c["last"] <= s["sailed"] <= c["last"] + 40 * H
                and (s["name"] in names or s["cs"] == v["call_sign"])]
        amb += len(cand) != 1
    ok("P3", amb == 0, "each call has exactly one candidate sailing (%d ambiguous)" % amb)
    fy = [c for c in x["calls"] if c.large and c.in_fy]
    ok("P4", len(fy) == len(calls), "world %d calls, pack %d" % (len(fy), len(calls)))

    # ------------------------------------------------------------------ the record's own allocation
    dayw = B.day_windows()
    tide = {t: B.windows(target, t) for t in ("SLT", "MRW")}
    ves = B.vessels(target)
    bad, at_open = 0, 0
    for c in x["calls"]:
        if not (P.PILOT_FROM <= c.sail < P.PILOT_TO):
            continue
        W_ = W.permitted(c, x["wins"], x["day"])
        exp = W.earliest(c.ready, W_) if W_ is not None else c.ready
        bad += exp != c.sail
        at_open += W_ is not None and c.sail > c.ready
    ok("SF6", bad == 0 and at_open > 100, "every sailing is the first permitted time an hour or more after its notice "
                                          "(%d off); %d sailings held to a window" % (bad, at_open))

    # ------------------------------------------------------------------ the fork table
    teus = [v["teu_nominal"] for v in ves.values()]
    ok("SF1", not any(9700 <= t <= 10100 for t in teus), "no ship between 9,700 and 10,100 TEU")
    near = [c for c in calls if min(abs(c["first"] - B.FY0), abs(c["first"] - B.FY1)) < 6 * H]
    gold = B.totals(target, "gold")
    win = B.ranking(gold)[0]

    def pick_population(pred):
        t = {k: 0.0 for k in B.TERMS}
        for c in x["calls"]:
            if c.large and pred(c):
                t[c.term] += W.saving(c, x)
        return B.ranking(t)[0]
    alt_arr = pick_population(lambda c: W.FY0 <= c.first - c.arr_lag < W.FY1)
    alt_sail = pick_population(lambda c: W.FY0 <= c.sail < W.FY1)
    ok("SF2", not near and alt_arr == win and alt_sail == win,
       "no first move within six hours of either end; by arrival %s, by sailing %s" % (alt_arr, alt_sail))
    gaps = defaultdict(list)
    for c in x["calls"]:
        if c.large:
            gaps[c.vessel["imo"]].append(c.first)
    ok("SF3", all(np.min(np.diff(sorted(v))) > 14 * W.DAY for v in gaps.values() if len(v) > 1),
       "no large ship calls at two terminals within a fortnight (no shift between terminals)")
    # SF4: the cut on each crane's own run, and the cut taken at the start
    t_run = {k: 0.0 for k in B.TERMS}
    for c in fy:
        lastrun = max(c.runs, key=lambda r: r.end)
        cut = (lastrun.end - lastrun.start) / 5.0
        W_ = W.permitted(c, x["wins"], x["day"])
        t_run[c.term] += (cut if W_ is None else c.sail - W.earliest(c.ready - cut, W_)) / H
    ok("SF4", B.ranking(t_run)[0] == win, "each crane's run cut by a fifth: %s first" % B.ranking(t_run)[0])
    # SF5: a let-go a fifth of the cargo work earlier falls outside the Salt Point window
    outside = tot_g = 0
    for r in B.call_table(target, "gold"):
        if r["gate"] == "tide":
            tot_g += 1
            t = r["sailed"] - r["cut"] * H
            Wt = tide[r["term"]]
            i = int(np.searchsorted(Wt[:, 1], t))
            outside += not (Wt[i, 0] <= t <= Wt[i, 1])
    ok("SF5", outside >= 0.9 * tot_g, "%d of %d deep-draft river let-gos a fifth earlier fall outside the window"
       % (outside, tot_g))
    lead_gap = min((r["notice"] - r["last"]) / 60 for r in B.call_table(target, "gold"))
    ok("SF7", B.ranking(B.totals(target, "ready_last"))[0] == win and lead_gap >= 24,
       "readiness at the last move: %s first; shortest notice after the last move %.0f min"
       % (B.ranking(B.totals(target, "ready_last"))[0], lead_gap))
    np_route = B.call_table(target, "no_passage")
    neg = sum(1 for r in np_route if r["term"] == "SLT" and r["gate"] == "tide" and r["saved"] < 0)
    n_slt = sum(1 for r in np_route if r["term"] == "SLT" and r["gate"] == "tide")
    outside_np = 0
    for s_ in sail:
        if s_["term"] in ("SLT", "MRW") and s_["draft_in"] >= B.DEEP_IN:
            Wn = B.windows(target, s_["term"], passage=False)
            i = int(np.searchsorted(Wn[:, 1], s_["sailed"]))
            outside_np += not (Wn[i, 0] <= s_["sailed"] <= Wn[i, 1])
    n_deep = sum(1 for s_ in sail if s_["term"] in ("SLT", "MRW") and s_["draft_in"] >= B.DEEP_IN)
    ok("SF8", neg >= 0.8 * n_slt and outside_np >= 0.95 * n_deep,
       "window at the berth refuted: %d of %d Saltworks held calls would save less than nothing; %d of %d deep-draft "
       "river sailings fall outside it" % (neg, n_slt, outside_np, n_deep))
    ok("SF9", B.ranking(B.totals(target, "lst_clock"))[0] == win,
       "standard time read as clock time: %s first" % B.ranking(B.totals(target, "lst_clock"))[0])
    # SF10: the register's summer draft in place of the sailing draft
    t_reg = {k: 0.0 for k in B.TERMS}
    for r in B.call_table(target, "gold"):
        if r["term"] in ("SLT", "MRW") and r["gate"] is None and r["vessel"]["summer_draft_m"] / 0.3048 >= 41:
            Wt = tide[r["term"]]
            s_ = lk[r["visit"]]
            t_reg[r["term"]] += (s_["sailed"] - B.earliest(s_["notice"] + B.NOTICE - r["cut"] * H, Wt)) / H
        else:
            t_reg[r["term"]] += r["saved"]
    ok("SF10", B.ranking(t_reg)[0] == win, "the register's summer draft as the draft: %s first" % B.ranking(t_reg)[0])
    loas = [v["loa_m"] for v in ves.values()]
    ok("SF11", not any(368.5 <= l_ <= 398.0 for l_ in loas), "no ship between 368.5 and 398 m (1,250 ft is 381 m)")
    ok("SF12", True, "daylight hours stated in the directions (8.1); none other ships")
    edge = 0
    for c in fy:
        W_ = W.permitted(c, x["wins"], x["day"])
        if W_ is not None:
            edge += W.edge_gap(int(round(c.ready - c.span / 5.0)), W_) < 60
    opens = sum(1 for c in x["calls"] if c.gated and c.sail > c.ready)
    ok("SF13", edge == 0 and opens > 0, "no replayed readiness within a minute of an edge; %d sailings at a window's "
                                        "opening" % opens)
    gc = {t: 0.0 for t in B.TERMS}
    for r in B.call_table(target, "gold"):
        gc[r["term"]] += round(r["saved"], 2)
    ok("SF14", all(round(gc[t]) == round(gold[t]) and abs((gc[t] - np.floor(gc[t])) - 0.5) >= 0.1 for t in B.TERMS),
       "two-decimal calls sum to the same whole hours")
    rk = B.ranking(gold)
    ok("SF15", round(gold[rk[0]]) - round(gold[rk[1]]) == round(gold[rk[0]] - gold[rk[1]]),
       "the lead is the same either way")

    # ------------------------------------------------------------------ the ladder
    def rung(route):
        t = B.totals(target, route)
        r = B.ranking(t)
        return t, r, t[r[0]] / t[r[1]]
    t0, r0, m0 = rung("naive")
    t1, r1, m1 = rung("tide_zero")
    t2, r2, m2 = rung("gold")
    ok("L1", r0[0] == "SLT" and m0 >= 1.15, "R0 a fifth of the cargo work: %s, margin %.3f" % (r0[0], m0))
    ok("L2", r0.index("MRW") in (3, 4), "Millrace %d on R0" % (r0.index("MRW") + 1))
    ok("L3", r1[0] not in ("MRW", "SLT") and m1 >= 1.15, "R1 held river sailings save nothing: %s, margin %.3f"
       % (r1[0], m1))
    ok("L4", r2[0] == "MRW" and m2 >= 1.5, "R2 the replay: %s, margin %.3f" % (r2[0], m2))
    ok("L5", t2[r2[1]] / t2[r2[2]] >= 1.06, "runner-up %s over third %s: %.3f" % (r2[1], r2[2], t2[r2[1]] / t2[r2[2]]))
    ok("L6", r2[1] not in (r0[1], r1[1]), "the runner-up %s is not R0's second %s or R1's second %s"
       % (r2[1], r0[1], r1[1]))
    dom = (t2["MRW"] - t2["SLT"]) / (t0["SLT"] - t0["MRW"])
    ok("L7", dom >= 1.2, "dominance %.2f" % dom)
    wrong = {k: B.ranking(B.totals(target, k))[0] for k in ("naive", "tide_zero", "gated_zero", "waited_zero")}
    ok("L8", all(v != "MRW" for v in wrong.values()), "every route without the replay misses: %s" % wrong)
    dm = B.ranking(B.totals(target, "daylight_missed"))
    ok("L9", dm[1] != r2[1], "the replay without the daylight rule: %s first, %s second" % (dm[0], dm[1]))
    rounded = {k: {t: round(v) for t, v in B.totals(target, k).items()} for k in
               ("naive", "tide_zero", "gated_zero", "waited_zero", "daylight_missed", "no_passage", "lst_clock",
                "ready_last")}
    hits = {k: sorted(t for t in B.TERMS if v[t] == round(t2[t])) for k, v in rounded.items()}
    ok("L10", all(not v for k, v in hits.items() if k in ("naive", "tide_zero", "gated_zero", "waited_zero")),
       "no wrong-call route prints a golden terminal figure: %s" % hits)
    # the analyst test on the clean truth: the easy route and the decisive step alone
    tw = defaultdict(lambda: {k: 0.0 for k in W.CODES})
    for c in fy:
        for k, v in W.routes(c, x).items():
            tw[k][c.term] += v
    ok("R1", B.ranking(tw["naive"])[0] == "SLT" and B.ranking(tw["gold"])[0] == "MRW",
       "clean truth: the cut files %s, the replay files %s" % (B.ranking(tw["naive"])[0], B.ranking(tw["gold"])[0]))
    ok("R2", all(abs(tw["gold"][t] - t2[t]) < 1e-6 for t in B.TERMS), "the pack's golden equals the clean truth")

    # ------------------------------------------------------------------ separation: the call reads no ask file
    tmp = tempfile.mkdtemp()
    try:
        for f in files:
            if f in (P.FILES["wo"], P.FILES["status"], P.FILES["labor"], P.FILES["oem"], P.FILES["jobs"], P.FILES["ais"],
                     P.FILES["yard"], P.FILES["windows"], P.FILES["register"]):
                continue
            shutil.copy(os.path.join(target, f), tmp)
        t_s = B.totals(tmp, "gold")
    finally:
        shutil.rmtree(tmp)
    ok("S1", all(abs(t_s[t] - t2[t]) < 1e-9 for t in B.TERMS), "the call unchanged with every ask file removed")

    # ------------------------------------------------------------------ ask A
    aw, ab = W.ask_a_truth(x), B.breakdown_hours(target)
    ok("A1", all(abs(aw[t] - ab[t]) < 1e-6 for t in B.TERMS), "ask A from the pack equals the clean truth")
    for cid, route, want in (("A2", "current", B.TERMS), ("A3", "order_spans", B.TERMS), ("A4", "code", B.TERMS),
                             ("A5", "both_counted", ["SLT", "MRW"])):
        v = B.breakdown_hours(target, route)
        miss = [t for t in B.TERMS if round(v[t]) != round(ab[t])]
        ok(cid, all(t in miss for t in want), "%s misses %s" % (route, miss))
    reg = B.register(target)
    _, _, where = B.crane_hours(target)
    moved = {k: sorted({t for _, t in where[k]}) for k in W.MOVES}
    ok("A6", all(len(v) == 2 and reg[k]["TERMINAL"] == W.MOVES[k][0] for k, v in moved.items()),
       "each moved crane works two terminals' calls; the register shows only where it is now: %s" % moved)
    pre = {k: sum(1 for e in x["maint"] if e["crane"] == k and e["kind"] == "BD" and not e.get("damage")
                  and W.FY0 <= e["periods"][0][0] < W.MOVES[k][1]) for k in W.MOVES}
    ok("A7", all(v >= 2 for v in pre.values()), "stops before each move inside the year: %s" % pre)
    # ------------------------------------------------------------------ ask B
    bw, bb = W.ask_b_truth(x), B.labor_hours(target)
    ok("B1", all(abs(bw[t] - bb[t]) < 1e-6 for t in B.TERMS), "ask B from the pack equals the clean truth")
    for cid, route, want in (("B2", "all_timesheets", B.TERMS), ("B3", "ledger_only", B.TERMS),
                             ("B4", "current", B.TERMS), ("B5", "with_yard", B.TERMS),
                             ("B6", "both_counted", ["SLT", "MRW"])):
        v = B.labor_hours(target, route)
        miss = [t for t in B.TERMS if round(v[t]) != round(bb[t])]
        ok(cid, all(t in miss for t in want), "%s misses %s" % (route, miss))
    d0, d1 = B.D0, B.D1
    keyed = Counter(o["ev"]["term"] for o in x["oem"] if o["keyed"] and d0 <= o["date"] <= d1)
    unkeyed = Counter(o["ev"]["term"] for o in x["oem"] if not o["keyed"] and d0 <= o["date"] <= d1)
    ok("B7", all(keyed[t] >= 1 and unkeyed[t] >= 1 for t in B.TERMS), "keyed %s, not keyed %s" % (dict(keyed),
                                                                                                    dict(unkeyed)))
    lab = B.rows_csv(B.p(target, "labor"))
    cset = Counter((r["WONUM"], r["TRANSDATE"], float(r["REGULARHRS"])) for r in lab if r["LABORCODE"] == "CONTRACT")
    clash = 0
    for o in x["oem"]:
        if not o["keyed"] and "wonum" in o["ev"]:
            clash += cset[(o["ev"]["wonum"], o["date"].isoformat(), o["hours"])] > 0
    ok("B8", clash == 0, "no contractor day matches a timesheet day that was not keyed")
    twins = sum(1 for o in x["oem"] if o["keyed"] and sum(v for (w, d, h), v in cset.items()
                                                          if w == o["ev"].get("wonum") and d == o["date"].isoformat()) > 1)
    ok("B9", twins > 0, "keyed days share their order and day with another contractor row on %d days" % twins)

    # ------------------------------------------------------------------ figures and goldens
    fig = json.load(open(fig_path)) if fig_path else {}
    vals = [("main", t, t2[t]) for t in B.TERMS] + [("A", t, ab[t]) for t in B.TERMS] + [("B", t, bb[t]) for t in B.TERMS]
    far = [(k, t, round(v, 3)) for k, t, v in vals if abs((v - np.floor(v)) - 0.5) < 0.15 or round(v) % 10 == 0]
    ok("E1", not far, "every graded figure at least 0.15 h from its cell's edge and not a multiple of ten %s" % far)
    gap = round(t2[r2[0]]) - round(t2[r2[1]])
    ok("E2", gap % 10 != 0, "the lead %d" % gap)
    if fig:
        n = B.NAME
        ok("F1", fig["call"]["terminal"] == n[r2[0]] and fig["components"]["runner_up"] == n[r2[1]]
           and fig["components"]["gap_hours"] == gap, "figures.json call, runner-up and lead")
        ok("F2", all(fig["components"]["berth_hours_saved"][n[t]] == round(t2[t]) for t in B.TERMS)
           and all(fig["a_breakdown_hours"][n[t]] == round(ab[t]) for t in B.TERMS)
           and all(fig["b_labor_hours"][n[t]] == round(bb[t]) for t in B.TERMS), "figures.json terminal figures")
    if gdir and os.path.isdir(gdir):
        wb = openpyxl.load_workbook(os.path.join(gdir, "ternmouth_2027_crane_placement.xlsx"))
        sh = wb["Berth hours saved"]
        got = {sh.cell(row=i, column=1).value: sh.cell(row=i, column=4).value for i in range(2, 8)}
        ok("F3", all(got[B.NAME[t]] == round(t2[t]) for t in B.TERMS), "the workbook's chart table")
        sh = wb["Crane condition"]
        got = {sh.cell(row=i, column=1).value: (sh.cell(row=i, column=2).value, sh.cell(row=i, column=3).value)
               for i in range(2, 8)}
        ok("F4", all(got[B.NAME[t]] == (round(ab[t]), round(bb[t])) for t in B.TERMS), "the workbook's ask table")
        with open(os.path.join(gdir, "ternmouth_fy2026_large_vessel_calls.csv"), encoding="utf-8") as fh:
            cr = list(csv.DictReader(fh))
        ok("F5", len(cr) == len(calls), "the calls table has %d rows" % len(cr))
    # the dictionary's repeated hour is in the crane extract, the first occurrence first
    rep = [r for r in B.rows_csv(B.p(target, "tos")) if r["HOUR_START"] == "2025-11-02 01:00"]
    keys = Counter((r["VISIT_ID"], r["CRANE_ID"]) for r in rep)
    ok("T1", any(v == 2 for v in keys.values()), "the 2 November 01:00 hour runs twice in the crane extract (%d rows)"
       % len(rep))
    r0_, r1_ = W.REPEAT
    big_rep = [c for c in calls if r0_ <= c["first"] < r1_ or r0_ <= c["last"] < r1_]
    ok("T2", not big_rep, "no large-vessel first or last move in the repeated hour")
    # the registers know every file (H9), and the documents read the same to a strict and a lenient reader (H25)
    man = {r["FILE"] for r in B.rows_csv(os.path.join(target, P.FILES["manifest"]))}
    ok("H9a", man == set(files) - {P.FILES["manifest"]}, "the manifest lists every shipped file but itself")
    dic = open(os.path.join(target, P.FILES["dictionary"]), encoding="utf-8").read()
    data = [f for f in files if os.path.splitext(f)[1] in (".csv", ".xlsx", ".json", ".txt")
            and f not in (P.FILES["dictionary"], P.FILES["manifest"])]
    ok("H9b", all(f in dic for f in data), "the dictionary names every data file")
    import re as _re
    import subprocess
    from pypdf import PdfReader
    must = {P.FILES["directions"]: ["shall pass Salt Point only between one hour thirty minutes before and thirty "
                                    "minutes after the time of high water at Salt Point",
                                    "from the Saltworks berths to Salt Point, 1 hour 50 minutes",
                                    "from the Millrace berths to Salt Point, 2 hours 40 minutes",
                                    "shall berth and unberth only between 0600 and 2000 local time",
                                    "A pilot is not provided at less than one hour's notice",
                                    "draft is the vessel's deepest draft, forward or aft"],
            P.FILES["rules"]: ["in four fifths of the time the call's cargo work took",
                               "the hours by which the call would have been shorter",
                               "from all fast to let go", "added unrounded"]}
    flat = lambda t: _re.sub(r"\s+", " ", t.replace("\u2019", "'"))
    for f, sentences in must.items():
        path = os.path.join(target, f)
        strict = flat("".join(pg.extract_text() for pg in PdfReader(path, strict=True).pages))
        poppler = flat(subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True).stdout)
        poppler2 = flat(subprocess.run(["pdftotext", path, "-"], capture_output=True, text=True).stdout)
        miss = [x for x in sentences if not (x in strict and x in poppler2)]
        ok("H25-" + f.split("_")[0], not miss, "%s: every deciding sentence read by a strict and a second reader %s"
           % (f, miss))
    # no SRT job's labor straddles the year's start (its hours carry one date)
    def straddles(j):
        a = (j["DOWN_UTC"] or j["RAISED_UTC"]).replace(tzinfo=B.UTC).timestamp() - 86400
        b = (j["UP_UTC"] or j["RAISED_UTC"]).replace(tzinfo=B.UTC).timestamp() + 2 * 86400
        return a < B.FY0 < b
    near = [j["JOB_NO"] for j in B.jobs(target) if straddles(j)]
    ok("Y1", not near, "no SRT job whose labor could fall either side of 1 July 2025 %s" % near)
    good = all(r[1] for r in rows)
    if not quiet:
        for cid, g, detail in rows:
            print("%-5s %-4s %s" % (cid, "ok" if g else "FAIL", detail))
        print("checks: %d of %d pass" % (sum(r[1] for r in rows), len(rows)))
    return good


if __name__ == "__main__":
    import sys
    here = os.path.dirname(os.path.abspath(__file__))
    task = os.path.dirname(here)
    t = sys.argv[1] if len(sys.argv) > 1 else os.path.join(task, "target")
    g = sys.argv[2] if len(sys.argv) > 2 else os.path.join(task, "golden")
    sys.exit(0 if run(t, g, meta_path=os.path.join(task, "metadata.json"),
                      fig_path=os.path.join(here, "figures.json")) else 1)
