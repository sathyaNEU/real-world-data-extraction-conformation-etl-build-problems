"""task242 documents: the city health officer's request (eml), the laboratory's cover note for the verification panel
(eml), the City Epidemiology Cell's indicator register (xlsx), SOP-EPI-07 on quality flags (pdf), the December dengue
bulletin (pdf), the serosurvey protocol (docx), the Vector Control Unit's surveillance SOP extract (pdf), the data
dictionary and the readme (md)."""
import datetime as dt

from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess
import records as RC

CORP = RC.CORP
PEOPLE = {"cho": ("Onveer Acharya", "cho"), "epi": ("Vedika Varghese", "vedika.varghese"),
          "micro": ("Shaurya Issac", "shaurya.issac"), "field": ("Gauri Kant", "gauri.kant")}
DOMAIN = "neithalpuramcorp.in"
COLLEGE = "ngmc.ac.in"


def _addr(k, dom=DOMAIN):
    name, user = PEOPLE[k]
    return "%s <%s@%s>" % (name, user, dom)


def _windows():
    out = []
    for z in P.ZONES:
        a = P.DAY0 + dt.timedelta(days=z[3] - P.FIELD_HALF)
        b = P.DAY0 + dt.timedelta(days=z[3] + P.FIELD_HALF)
        out.append((z[0], z[1], a, b))
    return out


def _span(a, b):
    if a.month == b.month:
        return "%d to %d %s" % (a.day, b.day, b.strftime("%b"))
    return "%d %s to %d %s" % (a.day, a.strftime("%b"), b.day, b.strftime("%b"))


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.4, leading=13, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=14, leading=18, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=14,
                        spaceBefore=8, spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10.5, textColor=colors.HexColor("#444444"))
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=8.2, leading=10)
    return body, h1, h2, small, cell


def _grid(rows, widths, cell, head=True):
    t = Table([[Paragraph(str(c), cell) for c in r] for r in rows], colWidths=widths, repeatRows=1 if head else 0)
    st = [("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA0AB")), ("VALIGN", (0, 0), (-1, -1), "TOP")]
    if head:
        st.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E2E7F0")))
    t.setStyle(TableStyle(st))
    return t


def _footer(text):
    def foot(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(20 * mm, 10 * mm, text)
        canvas.drawRightString(A4[0] - 20 * mm, 10 * mm, "Page %d" % d.page)
        canvas.restoreState()
    return foot


# ------------------------------------------------------------------ the request (eml)
def request(path, F):
    cho, epi, micro, field = _addr("cho"), _addr("epi"), _addr("micro", COLLEGE), _addr("field")
    body = f"""From: {cho}
To: {epi}
Cc: {field}
Subject: Serosurvey values to the health standing committee: quality flags
Date: Mon, 10 Jan 2028 08:41:37 +0530
Message-ID: <7f2c9e41b0a84d3e@mail.{DOMAIN}>
MIME-Version: 1.0
Content-Type: text/plain; charset="utf-8"
Content-Transfer-Encoding: 8bit

Vedika,

The public health standing committee meets on Monday 24 January to settle the 2028-29 vector control allocation by zone, and the zone tables go to the members on the 19th. The serosurvey values we published in the December bulletin (SERO-D1) have not been through the flag review yet, and SOP-EPI-07 says nothing we publish goes to a committee without it.

When the members see a zone's value I want its flag beside it, and in the same table the zone's Breteau index for the November round and the survey's household response rate in that zone, because they will ask about both. Please send me the workbook and one slide I can put up in the meeting.

Shaurya's verification run for the IgM kit lot the survey used is in the shared folder with his note; Gauri has put the field app export, the sample and the laboratory's specimen register there too, and the Vector Control Unit has sent the November round and the ward master.

Onveer Acharya
City Health Officer
{CORP}
"""
    with open(path, "w", encoding="utf-8", newline="\r\n") as fh:
        fh.write(body)


def panel_note(path, F):
    epi, micro = _addr("epi"), _addr("micro", COLLEGE)
    body = f"""From: {micro}
To: {epi}
Subject: IgM ELISA lot verification: results
Date: Tue, 4 Jan 2028 16:12:05 +0530
Message-ID: <a19d3c57e2f04b88@mail.{COLLEGE}>
MIME-Version: 1.0
Content-Type: text/plain; charset="utf-8"
Content-Transfer-Encoding: 8bit

Dear Vedika,

As you asked, the verification run we did in July for the IgM capture ELISA lot we have used since then, for routine work and for all the survey sera (file {F['panel']}).

The confirmed-case sera are from patients of our 2026 dengue follow-up clinic, confirmed by RT-PCR or NS1 ELISA and stored at minus 80; onset and collection dates are as the clinic recorded them. The controls are blood donors from March 2027 and patients with another confirmed febrile illness. As for any specimen, a first reading in the grey zone (index 0.90 to 1.10) was repeated and the repeat is the final result. All plates passed the kit's run validity criteria.

Regards,
Shaurya Issac
Department of Microbiology, Neithalpuram Government Medical College
"""
    with open(path, "w", encoding="utf-8", newline="\r\n") as fh:
        fh.write(body)


# ------------------------------------------------------------------ the indicator register (xlsx)
REGISTER = [
    ("DEN-01", "Dengue cases notified, by zone and week",
     "Cases in the city's dengue line list among residents of the zone, by week of notification",
     "Dengue cases among the zone's residents notified to the city's surveillance in the week",
     "City dengue line list", "Weekly", "Weekly dengue bulletin"),
    ("DEN-02", "Dengue notification rate",
     "DEN-01 for the year to date per 100,000 residents (2027 population estimate)",
     "The rate at which the zone's residents were notified as dengue cases in the year to date",
     "City dengue line list; population estimates", "Monthly", "Monthly dengue bulletin"),
    ("DEN-03", "Dengue deaths", "Deaths among notified residents with a dengue outcome of died",
     "Deaths from dengue among the zone's residents in the year", "City dengue line list", "Monthly",
     "Monthly dengue bulletin"),
    ("DEN-04", "Laboratory-confirmed share of notified cases",
     "Cases classified confirmed (ELISA or RT-PCR) as a per cent of DEN-01",
     "The share of notified dengue cases with laboratory confirmation", "City dengue line list", "Monthly",
     "Monthly dengue bulletin"),
    ("SERO-D1", "Dengue infection share, 2027 (household serosurvey)",
     "Residents with a reactive dengue IgM in the 2027 household serosurvey, per cent of residents tested",
     "The share of the zone's residents aged one year and over infected with dengue virus between 1 January 2027 "
     "and the collection of their survey specimen",
     "2027 household serosurvey; IgM capture ELISA, Department of Microbiology, NGMC", "Once, after the season",
     "December 2027 dengue bulletin"),
    ("CHK-01", "Chikungunya cases notified, by zone and week",
     "Cases in the city's chikungunya line list among residents of the zone, by week of notification",
     "Chikungunya cases among the zone's residents notified in the week", "City chikungunya line list", "Weekly",
     "Weekly fever bulletin"),
    ("MAL-01", "Malaria cases, by zone and month",
     "Smear or RDT positive malaria cases among residents of the zone, by month of diagnosis",
     "Malaria cases among the zone's residents diagnosed in the month", "Malaria register", "Monthly",
     "Monthly fever bulletin"),
    ("FEV-01", "Fever cases seen at UPHCs, by zone and week",
     "Fever cases recorded in the UPHC OPD registers of the zone, by week",
     "The fever load on the zone's primary health centres", "UPHC OPD registers", "Weekly",
     "Weekly fever bulletin"),
]


def register(path, F):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Register 2027"
    ws["A1"] = "City Epidemiology Cell: register of published indicators, 2027"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Quality flags on every published value are set under SOP-EPI-07 (rev. 3)."
    head = ["Code", "Indicator", "Published value", "What the value stands for", "Source", "Frequency", "Published in"]
    for j, h in enumerate(head, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE3EC")
    for i, row in enumerate(REGISTER, 5):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFG", [10, 30, 46, 52, 34, 14, 26]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    wb.save(path)
    postprocess.scrub(path, PEOPLE["epi"][0], "2027-04-03")
    postprocess.finish_office(path, "2027-04-01T10:05:00Z", "2027-12-29T11:32:00Z", CORP)


# ------------------------------------------------------------------ SOP-EPI-07 (pdf)
def sop(path, F):
    body, h1, h2, small, cell = _styles()
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                            bottomMargin=18 * mm, title="SOP-EPI-07 Quality flags on published surveillance indicators",
                            author="City Epidemiology Cell", creator=CORP, subject="Standard operating procedure",
                            invariant=1)
    st = [Paragraph("SOP-EPI-07: Quality flags on published surveillance indicators", h1),
          Paragraph("%s, Public Health Department, City Epidemiology Cell. Revision 3, effective 1 April 2027; "
                    "approved by the City Health Officer." % CORP, small), Spacer(1, 6),
          Paragraph("1. Purpose", h2),
          Paragraph("A published value is read by people who will not see how it was made. A quality flag tells them, "
                    "beside the value, that it should not be read as it stands for one of the reasons below.", body),
          Paragraph("2. Scope", h2),
          Paragraph("Every indicator in the Cell's register of published indicators, each time a value of it is "
                    "published or sent to a committee.", body),
          Paragraph("3. Flags", h2)]
    rows = [["Flag", "Applies to", "The flag goes on when"],
            ["P (provisional)", "Weekly counts", "the week is one of the three most recent, while late reports are "
                                                 "still arriving"],
            ["R (incomplete reporting)", "Weekly counts", "fewer than 80 per cent of the reporting units sent their "
                                                          "weekly report"],
            ["S (small numbers)", "Rates and shares", "the value rests on fewer than 10 events"],
            ["M (misclassification)", "Values computed from the results of a diagnostic test",
             "the quantity the register says the value stands for, estimated from the same data with allowance for "
             "the test's misclassification as the testing laboratory has measured it, is more than one and a half "
             "times the published value or less than two thirds of it"]]
    st += [_grid(rows, [36 * mm, 42 * mm, 92 * mm], cell), Spacer(1, 6),
           Paragraph("4. Procedure", h2),
           Paragraph("The Cell sets the flags before a value is released to a committee or to the public and keeps the "
                     "workings with the indicator. A flag travels with the value wherever it is reproduced. A value "
                     "that carries a flag is still published; the flag is not a correction of it.", body),
           Paragraph("5. Records", h2),
           Paragraph("The flag review for each release is kept as a workbook with the indicator's file, with each "
                     "value, the flag set and the figures the flag was decided on.", body)]
    doc.build(st, onFirstPage=_footer("%s | City Epidemiology Cell | SOP-EPI-07 rev. 3" % CORP),
              onLaterPages=_footer("%s | City Epidemiology Cell | SOP-EPI-07 rev. 3" % CORP))
    postprocess.scrub(path, "City Epidemiology Cell", "2027-03-28")


# ------------------------------------------------------------------ the December bulletin (pdf)
def bulletin(path, F, R):
    body, h1, h2, small, cell = _styles()
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title="City Dengue Bulletin No. 52 of 2027",
                            author="City Epidemiology Cell", creator=CORP, subject="Dengue surveillance", invariant=1)
    pub = R.published()
    win = {code: (a, b) for code, name, a, b in _windows()}
    c = R.cases_df
    res = c[(c["zone_k"] >= 0) & (c["notified_day"] >= 3) & (c["notified_day"] <= 359)]
    weeks = res["notified_day"].map(lambda d: (P.DAY0 + dt.timedelta(days=int(d))).isocalendar()[1]).value_counts()
    st = [Paragraph("City Dengue Bulletin", h1),
          Paragraph("%s, Public Health Department, City Epidemiology Cell. No. 52 of 2027, for the week ending Sunday "
                    "26 December 2027. Issued 28 December 2027." % CORP, small), Spacer(1, 6),
          Paragraph("1. The season", h2),
          Paragraph("Notified dengue cases among the city's residents rose from late July, peaked at the turn of "
                    "August and September, stayed high through September, fell in October and rose again in November "
                    "with the north-east monsoon, more steeply in some zones than in others. Figure 1 shows the "
                    "residents' cases by week of notification.", body)]
    d = Drawing(170 * mm, 62 * mm)
    bc = VerticalBarChart()
    bc.x, bc.y, bc.width, bc.height = 12 * mm, 10 * mm, 152 * mm, 46 * mm
    vals = [int(weeks.get(w, 0)) for w in range(1, 52)]
    bc.data = [vals]
    bc.bars[0].fillColor = colors.HexColor("#5B7DB1")
    bc.bars[0].strokeColor = None
    bc.valueAxis.valueMin = 0
    bc.valueAxis.valueMax = int((max(vals) // 200 + 1) * 200)
    bc.valueAxis.valueStep = 200
    bc.valueAxis.labels.fontSize = 6.5
    bc.categoryAxis.categoryNames = [str(w) if w % 4 == 1 else "" for w in range(1, 52)]
    bc.categoryAxis.labels.fontSize = 6.5
    d.add(bc)
    d.add(String(12 * mm, 2 * mm, "Week of notification, 2027", fontSize=7, fontName="Helvetica"))
    st += [d, Paragraph("Figure 1. Dengue cases notified among residents, by week of notification, weeks 1 to 51 of "
                        "2027 (week 52 provisional, not shown).", small),
           Paragraph("2. Household serosurvey 2027", h2),
           Paragraph("Between 13 October and 7 December the survey teams visited the 2,178 households sampled across "
                     "the 12 zones and took a blood specimen from each resident aged one year and over who agreed. The "
                     "Department of Microbiology, Neithalpuram Government Medical College, tested the sera for dengue "
                     "IgM. Table 1 gives each zone's dengue infection share for 2027 (indicator SERO-D1): the residents "
                     "whose serum was IgM reactive, as a per cent of the residents tested.", body)]
    rows = [["Zone", "Fieldwork", "Residents tested", "IgM reactive", "Dengue infection share, 2027 (%)"]]
    for _, x in pub.iterrows():
        a, b = win[x["code"]]
        rows.append(["%s %s" % (x["code"], x["zone"]), _span(a, b), "{:,}".format(int(x["tested"])),
                     str(int(x["reactive"])), "%.1f" % x["share"]])
    tot_n, tot_a = int(pub["tested"].sum()), int(pub["reactive"].sum())
    rows.append(["City", "13 Oct to 7 Dec", "{:,}".format(tot_n), str(tot_a), "%.1f" % (100.0 * tot_a / tot_n)])
    t = _grid(rows, [40 * mm, 34 * mm, 28 * mm, 24 * mm, 46 * mm], cell)
    st += [t, Paragraph("Table 1. Dengue infection share, 2027, by zone (SERO-D1). Source: 2027 household serosurvey; "
                        "IgM capture ELISA, Department of Microbiology, NGMC.", small),
           Paragraph("3. Vector control", h2),
           Paragraph("Fogging and source reduction drives continued in every zone through November. The Vector Control "
                     "Unit's entomological indices for the November round are reported to the public health standing "
                     "committee separately.", body)]
    doc.build(st, onFirstPage=_footer("%s | City Epidemiology Cell | Dengue bulletin 52/2027" % CORP),
              onLaterPages=_footer("%s | City Epidemiology Cell | Dengue bulletin 52/2027" % CORP))
    postprocess.scrub(path, "City Epidemiology Cell", "2027-12-28")


# ------------------------------------------------------------------ the serosurvey protocol (docx)
def protocol(path, F):
    import docx
    from docx.shared import Pt
    doc = docx.Document()
    stl = doc.styles["Normal"]
    stl.font.name = "Calibri"
    stl.font.size = Pt(10.5)

    def para(t, bold_lead=None, after=6):
        p = doc.add_paragraph()
        if bold_lead:
            r = p.add_run(bold_lead + " ")
            r.bold = True
        p.add_run(t)
        p.paragraph_format.space_after = Pt(after)
        return p

    h = doc.add_paragraph()
    r = h.add_run("Household dengue serosurvey 2027: protocol")
    r.bold = True
    r.font.size = Pt(13)
    for k, v in [("Version", "1.2, 24 September 2027"), ("Survey lead", "Dr Vedika Varghese, City Epidemiology Cell"),
                 ("Field coordinator", "Gauri Kant"),
                 ("Laboratory", "Department of Microbiology, Neithalpuram Government Medical College")]:
        para(v, k + ":", after=0)
    doc.add_paragraph()
    para("After the 2027 transmission season the Corporation will measure, zone by zone, how many residents carry "
         "dengue IgM antibodies, to inform the 2028-29 vector control and outbreak preparedness plans.", "1. Purpose.")
    para("In each of the 12 zones about 180 households are selected by systematic sampling from the 2027 household "
         "listing, with a random start. Every usual resident aged one year and over in a selected household is invited "
         "to give a venous blood specimen. Selected households are not replaced.", "2. Sample.")
    para("Two teams work in each zone over about ten days, in the order the field schedule sets. A team calls at each "
         "selected household; where nobody is at home it calls again on another day, up to three calls in all. Each "
         "call is recorded on the field app (KoboToolbox) with its outcome: 1 interview completed, 2 refused, 3 nobody "
         "at home, 4 vacant, 5 not a dwelling, 6 other.", "3. Fieldwork.")
    para("The team's phlebotomist draws 3 to 5 ml into a plain tube at the household, labels it with the specimen "
         "barcode TS-ZZ-HHH-L (zone, household serial in the zone, line number of the resident on the household "
         "roster) and keeps it in the team's cool box. Cool boxes reach the laboratory the next morning. The laboratory "
         "registers every specimen it receives and rejects any that is haemolysed, short or leaking; while the team is "
         "still working in the zone it calls again for a fresh specimen, labelled with the same barcode and the letter "
         "R.", "4. Specimens.")
    para("Sera are tested for dengue IgM by capture ELISA in plates of 88 specimens with the kit's controls; an index "
         "of 1.10 or more is reactive and 0.90 to 1.10 is tested again. The laboratory releases each result to the "
         "survey's results file.", "5. Laboratory.")
    para("The field app export, the laboratory's specimen register and the results file are kept by the City "
         "Epidemiology Cell. The household response rate is the share of sampled households, other than those the "
         "field team found vacant or not a dwelling, in which at least one member gave a blood specimen.",
         "6. Data and quality.")
    doc.core_properties.author = PEOPLE["epi"][0]
    doc.core_properties.title = "Household dengue serosurvey 2027: protocol"
    doc.save(path)
    postprocess.scrub(path, PEOPLE["epi"][0], "2027-09-24")
    postprocess.finish_office(path, "2027-09-02T07:40:00Z", "2027-09-24T12:18:00Z", CORP)


# ------------------------------------------------------------------ the Vector Control Unit's SOP extract (pdf)
def vc_sop(path, F):
    body, h1, h2, small, cell = _styles()
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                            bottomMargin=18 * mm, title="SOP-VC-02 Entomological surveillance (extract)",
                            author="Vector Control Unit", creator=CORP, subject="Standard operating procedure",
                            invariant=1)
    st = [Paragraph("SOP-VC-02: Entomological surveillance for Aedes (extract, sections 4 and 6)", h1),
          Paragraph("%s, Public Health Department, Vector Control Unit. Revision 5, March 2026." % CORP, small),
          Spacer(1, 6), Paragraph("4. The monthly house-to-house round", h2),
          Paragraph("Each ward's Domestic Breeding Checkers (DBCs) inspect every house on their ward list once in the "
                    "calendar month, indoors and outdoors, and record on the inspection form the number of containers "
                    "holding water and the number found with Aedes larvae or pupae. They empty, cover or treat every "
                    "positive container they find and record the source reduction. A house found locked is called on "
                    "again within the week; the call is recorded as LOCKED. A positive house is visited again about a "
                    "week later to check the source reduction.",
                    body),
          Paragraph("6. Indices", h2)]
    rows = [["Index", "Definition"],
            ["House index (HI)", "Houses with at least one positive container, per 100 houses inspected"],
            ["Container index (CI)", "Positive containers per 100 containers holding water"],
            ["Breteau index (BI)", "Positive containers per 100 houses inspected"]]
    st += [_grid(rows, [40 * mm, 130 * mm], cell), Spacer(1, 6),
           Paragraph("A month's indices are computed for each ward and each zone from the round's inspections: each "
                     "house inspected in the month counts once, with what was found at its round inspection. A locked "
                     "call and a follow-up visit are not round inspections.", body)]
    doc.build(st, onFirstPage=_footer("%s | Vector Control Unit | SOP-VC-02 extract" % CORP),
              onLaterPages=_footer("%s | Vector Control Unit | SOP-VC-02 extract" % CORP))
    postprocess.scrub(path, "Vector Control Unit", "2026-03-16")


# ------------------------------------------------------------------ the data dictionary and readme (md)
def dictionary(F):
    return f"""# Data dictionary

Dates are ISO (YYYY-MM-DD) unless a column says otherwise. Times in the field app export carry their offset; the
laboratory's register is in Indian Standard Time.

## {F['cases']}
The city's dengue line list as kept by the City Epidemiology Cell: every dengue case notified to the Corporation from 1
October 2026, extracted on 8 January 2028. One row per notification.

| Column | Meaning |
|---|---|
| case_no | Notification number (NCMC/DEN/year of notification/serial) |
| notified_on | Date the case was notified (DD-MM-YYYY) |
| onset_date | Date of onset of fever, as reported (DD-MM-YYYY) |
| age_yrs, sex | Age in completed years; M or F |
| zone | Zone of residence as the reporting unit wrote it; non-residents as Non-resident with their district |
| reporting_unit | The hospital, primary health centre or laboratory that notified the case |
| test_done | The test that supported the notification |
| classification | CONFIRMED (ELISA or RT-PCR) or PROBABLE (rapid test) |
| hospitalised | Y or N |
| outcome | RECOVERED, DIED, LAMA (left against medical advice), UNDER TREATMENT |

## {F['results']}
The household serosurvey's IgM results: one row per resident whose specimen the laboratory tested.

| Column | Meaning |
|---|---|
| specimen_id | Specimen barcode (TS-ZZ-HHH-L, with R for a specimen collected again) |
| household_id | ZZ-HHH: zone number and household serial in the zone |
| zone | Zone |
| member_line | The resident's line on the household roster |
| age_yrs, sex | Age in completed years; F or M |
| collected_on | Date the specimen was collected |
| plate | ELISA plate |
| igm_index | Final IgM index (sample to cut-off ratio) |
| igm_result | REACTIVE or NON-REACTIVE |

## {F['panel']}
The laboratory's verification run for the IgM capture ELISA lot used for the survey (July 2027). One row per serum.

| Column | Meaning |
|---|---|
| panel_no | Serial in the run |
| arm | CONFIRMED CASE or NEGATIVE CONTROL |
| patient_ref | Follow-up clinic number (FUC), blood bank number (BB) or febrile illness register number (OFI) |
| confirmed_by | How the case was confirmed, or what the control is |
| onset_date | Date of onset of illness (DD-MM-YYYY); blank for blood donors |
| collected_on | Date the serum was collected (DD-MM-YYYY) |
| index_first, reading_first | First IgM index and reading |
| index_repeat | Index on repeat testing, for a first reading in the grey zone |
| result_final | REACTIVE or NON-REACTIVE |

## {F['sample']}
The households sampled for the survey, one row each.

| Column | Meaning |
|---|---|
| household_id | ZZ-HHH |
| zone | Zone |
| listing_sl_no | Serial in the zone's sample |
| listing_address | Address as listed |

## {F['visits']}
The field app export (KoboToolbox): one row per submission of the household call form.

| Column | Meaning |
|---|---|
| _id, _uuid | Submission identifiers given by the server |
| start, end | When the form was opened and closed on the tablet (with offset) |
| today | Date on the tablet |
| deviceid | Tablet |
| zone, household_id | The household called on |
| visit_number | Call number entered by the interviewer |
| visit_outcome | 1 interview completed, 2 refused, 3 nobody at home, 4 vacant, 5 not a dwelling, 6 other |
| members_usual, members_eligible | Usual residents; residents aged one year and over (completed interviews) |
| interviewer | Interviewer code |
| __version__ | Form version |
| _submission_time | When the server received the submission (UTC) |

## {F['lab']}
The laboratory's register of the survey specimens it received, one row per specimen.

| Column | Meaning |
|---|---|
| lab_no | Laboratory accession number |
| barcode | Specimen barcode from the field |
| zone | Zone, from the barcode |
| received_on | Date and time received |
| condition | ACCEPTED, HAEMOLYSED, INSUFFICIENT |
| action | TESTED, REJECTED, REJECTED - RECOLLECTION REQUESTED |
| plate | ELISA plate, for a tested specimen |
| released_on | Date the result was released |

## {F['insp']}
The Domestic Breeding Checkers' house inspection forms for November 2027, one row per form.

| Column | Meaning |
|---|---|
| record_id | Form number |
| inspected_on | Date of the visit |
| ward_no | Ward number written on the form |
| premise_id | Property assessment number (ZZ/WWW/NNNNN) |
| premise_type | Residential |
| wet_containers | Containers holding water |
| positive_containers | Containers with Aedes larvae or pupae |
| source_reduction | Y if positive containers were emptied, covered or treated |
| remarks | The checker's remark |
| inspector | DBC code |

## {F['wards']}
The Corporation's ward master, one row per ward.

| Column | Meaning |
|---|---|
| ward_no | Ward number |
| ward_name | Ward name |
| zone | Zone |
| former_ward_nos | The ward numbers the ward's area carried under the earlier ward map |

## {F['rain']}
Daily rainfall at the Corporation's rain gauges, 1 October 2026 to 31 December 2027.

| Column | Meaning |
|---|---|
| station_id, station_name | Gauge |
| date | Day |
| rainfall_mm | Rainfall in the 24 hours to 08:30 IST |

## {F['fog']}
The Vector Control Unit's log of fogging and source reduction operations in 2027.

| Column | Meaning |
|---|---|
| date | Day |
| zone, locality | Where |
| activity | Operation |
| houses_covered | Houses covered (not recorded for thermal fogging) |
| insecticide_litres | Insecticide used |
| trigger | Why the operation was run |
"""


def readme(F):
    return f"""# Serosurvey flag review, January 2028

Files gathered by the City Epidemiology Cell, {CORP}, for the quality flag review of the 2027 serosurvey values
before the public health standing committee.

- {F['request']}: the City Health Officer's request
- {F['register']}: the Cell's register of published indicators
- {F['standard']}: SOP-EPI-07, quality flags on published surveillance indicators
- {F['bulletin']}: City Dengue Bulletin No. 52 of 2027
- {F['panelnote']}: the laboratory's note on the kit lot's verification run
- {F['panel']}: the verification run's results
- {F['protocol']}: the household serosurvey protocol
- {F['results']}: the survey's IgM results
- {F['sample']}: the sampled households
- {F['visits']}: the field app export of household calls
- {F['lab']}: the laboratory's register of survey specimens
- {F['cases']}: the city's dengue line list
- {F['vcsop']}: the Vector Control Unit's entomological surveillance SOP (extract)
- {F['insp']}: the November 2027 house inspection forms
- {F['wards']}: the ward master
- {F['rain']}: daily rainfall at the Corporation's rain gauges
- {F['fog']}: the fogging and source reduction log
- DATA_DICTIONARY.md: the columns of every table
"""


def write_all(target, F, R):
    import os
    p = lambda k: os.path.join(target, F[k])
    request(p("request"), F)
    panel_note(p("panelnote"), F)
    register(p("register"), F)
    sop(p("standard"), F)
    bulletin(p("bulletin"), F, R)
    protocol(p("protocol"), F)
    vc_sop(p("vcsop"), F)
    with open(p("dict"), "w") as fh:
        fh.write(dictionary(F))
    with open(p("readme"), "w") as fh:
        fh.write(readme(F))
