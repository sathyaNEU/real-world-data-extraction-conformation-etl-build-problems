"""task242 goldens: the flag review workbook the City Epidemiology Cell keeps with the indicator (SOP-EPI-07 section 5)
and the slide for the public health standing committee."""
import io
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import params as P
import postprocess

XLSX = "dengue_igm_flag_review_2027.xlsx"
PPTX = "dengue_igm_flags_health_committee.pptx"
CORP = "Neithalpuram City Municipal Corporation"
AUTHOR = "Vedika Varghese"
NAVY, RED, GREY, LIGHT = "1F3A5F", "B03A2E", "8A8F98", "DDE3EC"


def _head(ws, title, sub):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=12, color=NAVY)
    ws["A2"] = sub
    ws["A2"].font = Font(size=9, color="555555")


def _table(ws, row, header, rows, widths, formats, flag_col=None):
    thin = Side(style="thin", color="B7BEC9")
    for j, h in enumerate(header, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=LIGHT)
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
        c.border = Border(bottom=thin)
    for i, r in enumerate(rows, row + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            if formats[j - 1]:
                c.number_format = formats[j - 1]
                c.alignment = Alignment(horizontal="right")
            if flag_col and j == flag_col and v == "M":
                c.font = Font(bold=True, color=RED)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + j)].width = w
    return row + len(rows)


def workbook(path, A):
    wb = Workbook()
    ws = wb.active
    ws.title = "Flag review"
    _head(ws, "SERO-D1 dengue infection share, 2027: misclassification flag review",
          "City Epidemiology Cell, %s. Review under SOP-EPI-07 for the public health standing committee of 24 January "
          "2028; prepared 12 January 2028." % CORP)
    hdr = ["Zone", "Fieldwork", "Residents tested", "IgM reactive", "Published share, 2027 (%)",
           "Estimated share infected since 1 January (%)", "Estimate / published", "Flag",
           "Breteau index, November 2027 round", "Household response rate (%)"]
    rows = []
    for z in A["zones"]:
        rows.append(["%s %s" % (z["code"], z["name"]), "%s to %s" % (z["start"][5:].replace("-", "/"), z["end"][5:].replace("-", "/")),
                     z["n"], z["a"], z["pub"], z["est"], z["ratio"], "M" if z["flag"] else "", z["bi"], z["rr"]])
    end = _table(ws, 4, hdr, rows, [20, 14, 11, 10, 13, 17, 11, 7, 15, 13],
                 [None, None, "#,##0", "#,##0", "0.0", "0.0", "0.00", None, "0.0", "0.0"], flag_col=8)
    ws.cell(row=end + 2, column=1, value="M: misclassification flag (SOP-EPI-07, section 3): the estimated share is more "
                                         "than 1.5 times the published value or less than two thirds of it.").font = Font(size=9)
    ws.cell(row=end + 3, column=1, value="Fieldwork dates are month/day, 2027. Estimated shares are accepted within %d per "
                                         "cent; smoothing the panel's curve another reasonable way moves none by more."
                                         % int(round(100 * A["band"]))).font = Font(size=9)
    ws.freeze_panes = "B5"
    ws.row_dimensions[4].height = 58
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1

    ws2 = wb.create_sheet("Assay sensitivity")
    _head(ws2, "IgM capture ELISA lot verification: sensitivity by days since onset",
          "Department of Microbiology, NGMC, July 2027 run; confirmed-case sera from the 2026 follow-up clinic, final "
          "readings. Specificity from %d control sera." % A["controls"])
    rows = [[g["visit"], "%d to %d" % (g["lo"], g["hi"]), g["mean_day"], g["n"], g["pos"], g["se"]] for g in A["groups"]]
    end = _table(ws2, 4, ["Scheduled visit (day)", "Days since onset", "Mean day", "Sera", "Reactive", "Sensitivity"],
                 rows, [14, 14, 10, 8, 9, 12], [None, None, "0.0", "#,##0", "#,##0", "0.000"])
    ws2.cell(row=end + 2, column=1, value="Specificity").font = Font(bold=True)
    ws2.cell(row=end + 2, column=6, value=A["sp"]).number_format = "0.0000"
    ws2.cell(row=end + 3, column=1, value="Control sera reactive: %d of %d" % (A["controls_reactive"], A["controls"]))
    rows = [["%s %s" % (z["code"], z["name"]), z["se"]] for z in A["zones"]]
    _table(ws2, end + 6, ["Zone", "Effective sensitivity for 2027 infections"], rows, [20, 18, 10, 8, 9, 12],
           [None, "0.000"])

    ws3 = wb.create_sheet("Notes")
    notes = [
        ("Value reviewed", "SERO-D1, dengue infection share 2027: residents with a reactive dengue IgM, per cent of residents "
                           "tested (City Dengue Bulletin No. 52 of 2027, Table 1)."),
        ("What it stands for", "The share of the zone's residents aged one year and over infected with dengue virus between "
                               "1 January 2027 and the collection of their survey specimen (indicator register)."),
        ("Estimate", "Rogan-Gladen: (p + Sp - 1) / (Se + Sp - 1). Se for a zone is the panel's sensitivity by days since "
                     "onset averaged over the zone's 2027 cases by onset date up to each specimen's collection, then over "
                     "the zone's tested residents; Sp from the panel's control sera."),
        ("Epidemic curves", "City dengue line list extracted 8 January 2028: residents' cases with onset from 1 January "
                            "2027, zone spellings mapped to the 12 zones, non-residents left out."),
        ("Breteau index", "Positive containers per 100 houses inspected in the November 2027 round (SOP-VC-02): each house "
                          "once, at its first inspection in the month; locked calls and follow-up visits are not round "
                          "inspections; zone from the premise's assessment number."),
        ("Response rate", "Sampled households in which at least one member gave a blood specimen, per cent of sampled "
                          "households other than those found vacant or not a dwelling (survey protocol, section 6); a "
                          "specimen the laboratory received counts whether it was tested or rejected."),
        ("As of", "Files as received on 10 January 2028."),
    ]
    _head(ws3, "Notes", "Definitions, sources and conventions used in this review.")
    for i, (k, v) in enumerate(notes, 4):
        ws3.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = ws3.cell(row=i, column=2, value=v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws3.row_dimensions[i].height = 44
    ws3.column_dimensions["A"].width = 20
    ws3.column_dimensions["B"].width = 110
    wb.save(path)
    postprocess.scrub(path, AUTHOR, "2028-01-12")
    postprocess.finish_office(path, "2028-01-11T05:20:00Z", "2028-01-12T09:46:00Z", CORP)


def chart_png(A):
    Z = A["zones"]
    fig, ax = plt.subplots(figsize=(11.2, 4.6), dpi=150)
    x = list(range(len(Z)))
    pub = [z["pub"] for z in Z]
    est = [z["est"] for z in Z]
    ax.bar([i - 0.19 for i in x], pub, width=0.36, color="#9FB3CF", label="Published share (IgM reactive)")
    ax.bar([i + 0.19 for i in x], est, width=0.36, color=["#%s" % RED if z["flag"] else "#%s" % NAVY for z in Z],
           label="Estimated share infected since 1 January")
    import matplotlib.patheffects as pe
    halo = [pe.Stroke(linewidth=3.2, foreground="white"), pe.Normal()]
    for i, z in enumerate(Z):
        lo, hi = z["pub"] * P.TOL_LOW, z["pub"] * P.TOL_HIGH
        ax.plot([i + 0.01, i + 0.37], [hi, hi], color="#222222", lw=1.5, path_effects=halo)
        ax.plot([i + 0.01, i + 0.37], [lo, lo], color="#222222", lw=1.1, ls=(0, (2, 1.5)), path_effects=halo)
        if z["flag"]:
            ax.text(i + 0.19, z["est"] + 0.5, "M", ha="center", va="bottom", fontsize=10, fontweight="bold",
                    color="#%s" % RED)
    ax.plot([], [], color="#333333", lw=1.4, label="Flag limit: 1.5 x published")
    ax.plot([], [], color="#333333", lw=1.0, ls=(0, (2, 1.5)), label="Flag limit: two thirds of published")
    ax.set_xticks(x)
    ax.set_xticklabels(["%s\n%s" % (z["code"], z["name"]) for z in Z], fontsize=8)
    ax.set_ylabel("Per cent of residents")
    ax.set_ylim(0, max(est + [z["pub"] * P.TOL_HIGH for z in Z]) * 1.08)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.legend(fontsize=8, frameon=False, loc="upper center", bbox_to_anchor=(0.5, -0.16), ncol=4)
    ax.yaxis.grid(True, color="#E3E6EA", lw=0.6)
    ax.set_axisbelow(True)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def slide(path, A):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fl = [z for z in A["zones"] if z["flag"]]
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    words = {1: "One", 2: "Two", 3: "Three", 4: "Four", 5: "Five", 6: "Six", 7: "Seven", 8: "Eight", 9: "Nine"}
    p.text = ("%s of the 12 zones' published 2027 infection shares carry the misclassification flag: their IgM shares "
              "stand for under two thirds of the residents infected since January" % words.get(len(fl), str(len(fl))))
    p.runs[0].font.size, p.runs[0].font.bold = Pt(22), True
    p.runs[0].font.color.rgb = RGBColor.from_string(NAVY)
    p2 = tf.add_paragraph()
    p2.text = ("Flagged (M): %s. The other zones' values stand within the tolerance." %
               ", ".join("%s %s (%.1f%% published, %.1f%% estimated)" % (z["code"], z["name"], z["pub"], z["est"])
                         for z in fl))
    p2.runs[0].font.size = Pt(11)
    pic = chart_png(A)
    s.shapes.add_picture(io.BytesIO(pic), Inches(0.5), Inches(1.55), width=Inches(12.3))
    sb = s.shapes.add_textbox(Inches(0.5), Inches(6.85), Inches(11.3), Inches(0.4))
    q = sb.text_frame.paragraphs[0]
    q.text = ("Source: 2027 household serosurvey (SERO-D1, bulletin 52/2027); IgM ELISA lot verification, Department of "
              "Microbiology, NGMC; city dengue line list, 8 January 2028. Flag M under SOP-EPI-07. City Epidemiology Cell.")
    q.runs[0].font.size = Pt(9)
    q.runs[0].font.color.rgb = RGBColor.from_string("555555")
    nb = s.shapes.add_textbox(Inches(12.3), Inches(6.85), Inches(0.6), Inches(0.4))
    nb.text_frame.paragraphs[0].text = "1"
    nb.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    cp = prs.core_properties
    cp.author, cp.last_modified_by, cp.comments = AUTHOR, AUTHOR, ""
    cp.title = "SERO-D1 misclassification flags, 2027"
    prs.save(path)
    postprocess.scrub(path, AUTHOR, "2028-01-12")
    postprocess.finish_office(path, "2028-01-12T06:02:00Z", "2028-01-12T10:15:00Z", CORP)


def write(folder, A):
    os.makedirs(folder, exist_ok=True)
    workbook(os.path.join(folder, XLSX), A)
    slide(os.path.join(folder, PPTX), A)
