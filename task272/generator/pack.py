"""Readers of the cut pack, shared by the analysis and the cracker: every figure the build reports is computed from the
files in target/, never from the world. The three hubs' extracts are conformed to one row shape: hospital, local
datetime, user, application, guarantor, type (INITIAL, RECON, PRESUMPTIVE), outcome (approve, deny, pend, withdraw),
the percent written off, the raw reason code, and for an application its received date, household size, the main
earner's work (the Willamette hub's numbers) and the 2025 policy's yearly income figure in dollars (the Puget hub's
three months times four, the Willamette hub's own figure, the Columbia hub's monthly average times twelve)."""
import csv
import datetime as dt
import os
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

import common as K

PACIFIC = ZoneInfo("America/Los_Angeles")
SYSTEM = set(K.SYSTEM_USER.values())
PUGET_WORK = {lbl: n for _, n, _, lbl in K.WORK}
COLUMBIA_WORK = {s: n for _, n, s, _ in K.WORK}


def rows(target, key):
    with open(os.path.join(target, K.FILES[key]), encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def _num(s):
    return float(s) if s not in ("", None) else None


def conformed(target, utc_as_local=False):
    """Every row of the three extracts in one shape. utc_as_local reads the Columbia clock as if it were local."""
    out = []
    for r in rows(target, "puget"):
        inc = _num(r["income_3mo"])
        out.append(dict(code=r["hospital_code"], when=dt.datetime.strptime(r["determined_at"], "%Y-%m-%d %H:%M:%S"),
                        user=r["user_id"], app=r["application_id"], guarantor=r["guarantor_id"],
                        type={"INITIAL": "INITIAL", "RECONSIDERATION": "RECON",
                              "PRESUMPTIVE": "PRESUMPTIVE"}[r["determination_type"]],
                        outcome={"APPROVED": "approve", "DENIED": "deny", "PENDED": "pend",
                                 "WITHDRAWN": "withdraw"}[r["outcome"]],
                        pct=int(r["discount_pct"]) if r["discount_pct"] else None,
                        reason=int(r["denial_reason"]) if r["denial_reason"] else None, hub="puget",
                        rid=r["determination_id"],
                        received=dt.date.fromisoformat(r["received_date"]) if r["received_date"] else None,
                        size=int(r["household_size"]) if r["household_size"] else None,
                        annual=round(inc * 4, 2) if inc is not None else None,
                        work=PUGET_WORK.get(r["work_type"])))
    for r in rows(target, "willamette"):
        when = dt.datetime.strptime(r["DET_DATE"] + " " + r["DET_TIME"], "%m/%d/%Y %H:%M:%S")
        out.append(dict(code=r["HOSP_ID"], when=when, user=r["USER_ID"], app=r["APPL_NO"], guarantor=r["GUAR_NO"],
                        type="PRESUMPTIVE" if r["USER_ID"] in SYSTEM else ("RECON" if r["RECON"] == "Y" else "INITIAL"),
                        outcome={"A": "approve", "D": "deny", "P": "pend", "X": "withdraw"}[r["RESULT"]],
                        pct=int(r["ADJ_PCT"]) if r["ADJ_PCT"] else None,
                        reason=int(r["DENY_RSN"]) if r["DENY_RSN"] else None, hub="willamette", rid=r["DET_NO"],
                        received=dt.datetime.strptime(r["RCVD_DATE"], "%m/%d/%Y").date() if r["RCVD_DATE"] else None,
                        size=int(r["HH_SIZE"]) if r["HH_SIZE"] else None, annual=_num(r["ANN_INC"]),
                        work=int(r["WORK_CD"]) if r["WORK_CD"] else None))
    for r in rows(target, "columbia"):
        t = dt.datetime.strptime(r["ts_utc"], "%Y-%m-%dT%H:%M:%SZ")
        when = t if utc_as_local else t.replace(tzinfo=ZoneInfo("UTC")).astimezone(PACIFIC).replace(tzinfo=None)
        d = int(r["decision"])
        inc = _num(r["inc_mo"])
        out.append(dict(code=r["facility"], when=when, user=r["staff_key"], app=r["app_ref"], guarantor=r["guarantor"],
                        type={"new": "INITIAL", "recheck": "RECON", "screen": "PRESUMPTIVE"}[r["kind"]],
                        outcome={1: "approve", 2: "approve", 3: "deny", 8: "withdraw", 9: "pend"}[d],
                        pct=int(r["pct"]) if r["pct"] else None, reason=int(r["reason"]) if r["reason"] else None,
                        hub="columbia", rid=r["seq"],
                        received=dt.date.fromisoformat(r["rcvd"]) if r["rcvd"] else None,
                        size=int(r["hh"]) if r["hh"] else None,
                        annual=round(inc * 12, 2) if inc is not None else None,
                        work=COLUMBIA_WORK.get(r["work"])))
    return out


def regions(target):
    wb = load_workbook(os.path.join(target, K.FILES["hospitals"]), read_only=True)
    ws = wb["Members 2025"]
    head = None
    out = {}
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == "Code":
            head = list(row)
            continue
        if head and row and row[0]:
            out[row[0]] = row[head.index("Region")]
    return out


def guarantor_map(target):
    """H23's Puget guarantor numbers back to their Columbia numbers, so one household keeps one key."""
    wb = load_workbook(os.path.join(target, K.FILES["crosswalk"]), read_only=True)
    ws = wb["Guarantor numbers"]
    out = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row and row[0]:
            out[row[2]] = row[1]
    return out


def first_reviews(target, rows_=None):
    """Every counsellor's first approval or denial of an application, in either year, as the models read them."""
    rows_ = conformed(target) if rows_ is None else rows_
    reg = regions(target)
    gmap = guarantor_map(target)
    work_kind = {n: k for k, n, _, _ in K.WORK}
    out = []
    for r in rows_:
        if r["type"] != "INITIAL" or r["outcome"] not in ("approve", "deny") or r["user"] in SYSTEM:
            continue
        g = gmap.get(r["guarantor"], r["guarantor"]) if r["code"] == K.MIGRATED else r["guarantor"]
        out.append(dict(code=r["code"], region=reg[r["code"]], kind=work_kind[r["work"]], hh="%s|%s" % (r["code"], g),
                        size=r["size"], month=r["received"].month,
                        mindex=(r["received"].year - 2024) * 12 + r["received"].month - 1, annual=r["annual"],
                        year=r["when"].year, writeoff=r["pct"] if r["outcome"] == "approve" else 0, app=r["app"]))
    return out


def ledger(target, dedupe=True):
    L = rows(target, "ledger")
    if not dedupe:
        return L
    seen, out = set(), []
    for r in L:
        k = tuple(r.values())
        if k not in seen:
            seen.add(k)
            out.append(r)
    return out
