"""The goldens: the program team's workbook (each member hospital, its eligible households' first-review decisions,
those that left the household owing more than half, the share, whether it qualifies, and the two scorecard columns)
and the chart of the 31 shares against the six-percent line, plus figures.json."""
import decimal
import json
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt                       # noqa: E402
from openpyxl import Workbook                         # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side   # noqa: E402
from openpyxl.utils import get_column_letter          # noqa: E402

import analysis as A                                  # noqa: E402
import common as K                                    # noqa: E402
import docs as DC                                     # noqa: E402
import models as M                                    # noqa: E402

TEAL = "1D5C63"
THIN = Side(style="thin", color="B8BFC4")


def half_up(x, places):
    q = decimal.Decimal(1).scaleb(-places)
    return float(decimal.Decimal(repr(float(x))).quantize(q, rounding=decimal.ROUND_HALF_UP))


def counts(fr):
    """Each hospital's 2025 first-review decisions, those within the limit on the golden estimate, and those of them
    that left the household owing more than half its balance."""
    cals = M.cal_cross(fr, M.k_region_kind)
    est = M.by_calendar(cals, M.k_region_kind)
    n, e, o = ({c: 0 for c in K.CODES} for _ in range(3))
    for x in fr:
        if x["year"] != K.YEAR:
            continue
        n[x["code"]] += 1
        if M.pct(x, est(x)) <= M.LIMIT:
            e[x["code"]] += 1
            o[x["code"]] += x["writeoff"] < M.OWING
    return n, e, o


def figures(res):
    c = res["call"]
    fr = c["first_reviews"]
    fam = M.family(fr)
    ref = M.reflexes(fr)
    qualify = c["qualify"]
    n, e, o = counts(fr)
    a_k = {code: half_up(res["ask_a"][code] / 1000.0, 1) for code in K.CODES}
    margin = min((fam[k][code] - K.THRESHOLD) if code in qualify else (K.THRESHOLD - fam[k][code])
                 for code in K.CODES for k in fam)
    return {
        "call": {"count": len(qualify), "qualify": qualify, "qualify_names": [K.NAME[x] for x in qualify],
                 "not": [x for x in K.CODES if x not in qualify]},
        "budget_usd": len(qualify) * DC.GRANT,
        "grant_usd": DC.GRANT,
        "verdict": {x: ("Yes" if x in qualify else "No") for x in K.CODES},
        "measure": {x: round(float(c["measure"][x]), 4) for x in K.CODES},
        "measure_1dp": {x: half_up(c["measure"][x], 1) for x in K.CODES},
        "family": {k: {x: round(float(v[x]), 4) for x in K.CODES} for k, v in fam.items()},
        "family_margin_floor": round(float(margin), 3),
        "reflex": {k: {x: round(float(v.get(x, float("nan"))), 4) for x in K.CODES} for k, v in ref.items()},
        "reflex_counts": {k: sum(1 for x in K.CODES if ref[k].get(x, 0) >= K.THRESHOLD) for k in ref},
        "reflex_sets": {k: [x for x in K.CODES if ref[k].get(x, 0) >= K.THRESHOLD] for k in ref},
        "first_reviews": n,
        "first_reviews_total": sum(n.values()),
        "eligible": e,
        "eligible_total": sum(e.values()),
        "owing": o,
        "owing_total": sum(o.values()),
        "ask_a": a_k,
        "ask_a_usd": {x: round(res["ask_a"][x], 2) for x in K.CODES},
        "ask_a_total": half_up(sum(res["ask_a"].values()) / 1000.0, 1),
        "ask_b": dict(res["ask_b"]),
        "ask_b_total": sum(res["ask_b"].values()),
        "filed": {"threshold_percent": K.THRESHOLD, "year": K.YEAR,
                  "clauses": ["1.2", "2.1", "2.2", "2.3", "2.4", "2.5"], "limit_percent_of_guideline": int(M.LIMIT),
                  "owing_discount_percent": 25, "extract_cut": K.EXTRACT_CUT.isoformat()},
    }


def _head(ws, row, cols, widths):
    for j, (h, wd) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row, j, h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = wd


def workbook(path, fig):
    wb = Workbook()
    ws = wb.active
    ws.title = "Grants FY2027"
    n = fig["call"]["count"]
    ws["A1"] = "Decision Coaching Grants FY2027: %d grants, $%s" % (n, format(fig["budget_usd"], ","))
    ws["A1"].font = Font(bold=True, size=13, color=TEAL)
    ws["A2"] = ("%d of the collaborative's 31 hospitals qualify under the criteria approved 8 December 2026: %s."
                % (n, ", ".join(fig["call"]["qualify_names"])))
    ws["A2"].font = Font(italic=True, size=10)
    cols = ["Code", "Hospital", "State", "Region", "First-review decisions, 2025",
            "Of them, households within the free-care limit on twelve-month income",
            "Of those, left owing more than half the balance", "Share (%)", "Qualifies",
            "Assistance written off on counsellor approvals, net of reversals ($000)",
            "Counsellor denials for documents not received"]
    _head(ws, 4, cols, [7, 34, 7, 22, 13, 17, 15, 9, 10, 20, 16])
    ws.row_dimensions[4].height = 76
    r = 5
    for code in K.CODES:
        st = next(h[3] for h in K.HOSPITALS if h[0] == code)
        vals = [code, K.NAME[code], st, K.REGION[code], fig["first_reviews"][code], fig["eligible"][code],
                fig["owing"][code], fig["measure_1dp"][code], fig["verdict"][code], fig["ask_a"][code],
                fig["ask_b"][code]]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(r, j, v)
            cell.border = Border(bottom=THIN)
            if j in (5, 6, 7, 11):
                cell.number_format = "#,##0"
            elif j == 8:
                cell.number_format = "0.0"
            elif j == 10:
                cell.number_format = "#,##0.0"
            if j >= 5:
                cell.alignment = Alignment(horizontal="center" if j == 9 else "right")
        if fig["verdict"][code] == "Yes":
            for j in range(1, 12):
                ws.cell(r, j).fill = PatternFill("solid", fgColor="E3F0F1")
                ws.cell(r, j).font = Font(bold=True)
        r += 1
    ws.cell(r, 1, "All hospitals").font = Font(bold=True)
    for j, v, fmt in ((5, fig["first_reviews_total"], "#,##0"), (6, fig["eligible_total"], "#,##0"),
                      (7, fig["owing_total"], "#,##0"), (9, "%d of 31" % n, None),
                      (10, fig["ask_a_total"], "#,##0.0"), (11, fig["ask_b_total"], "#,##0")):
        cell = ws.cell(r, j, v)
        if fmt:
            cell.number_format = fmt
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center" if j == 9 else "right")
        cell.border = Border(top=Side(style="thin", color="1F2428"))
    ws.freeze_panes = "C5"
    nt = wb.create_sheet("Notes")
    notes = [
        "How the figures were made",
        "Qualifies (criteria 2.1 to 2.5): each hospital's 2025 first-review decisions by local date (the Columbia "
        "extract's UTC clock put on Pacific time; H23's rows from both hubs); reconsiderations, pends, withdrawals and "
        "the presumptive screening left out. The extracts hold each household's income only as the 2025 policy "
        "counted it, the three full months before the month of application times four, each hub in its own form "
        "(Puget the three months' total, Willamette the yearly figure, Columbia the monthly average). A household's "
        "income over the twelve months before it applied is that yearly figure divided by the factor of its month of "
        "application on the calendar of its kind of work in its region: the policy's figure averaged over every "
        "first-review decision of 2024 and 2025 in that region and kind of work received in that month, set against "
        "its average over the twelve months. Harvest, packing, fishing and resort work peak in different months in "
        "different regions; year-round work's factors stay close to one in every month.",
        "Check: the calendars read within households that applied in more than one month (a household effect and a "
        "month effect), on geometric means, from 2025's decisions alone, by hospital rather than region, or for the "
        "seasonal kinds of work alone put every hospital on the same side of six percent (the nearest is %.1f points "
        "from the line)." % fig["family_margin_floor"],
        "Grant: $%s for twelve months (criteria 1.2); FY2027 budget line %d x $%s = $%s." % (
            format(fig["grant_usd"], ","), fig["call"]["count"], format(fig["grant_usd"], ","),
            format(fig["budget_usd"], ",")),
        "Assistance written off: write-offs (FAWO) on approvals a counsellor recorded in 2025, first review or "
        "reconsideration, each write-off taken once (the adjustments extract repeats it on every account the approval "
        "covers) and credited to the household's approval decided last before it posted (a household approved on a "
        "later application for the same bill has its write-off posted under the number its case was opened with), "
        "less reversals (FAWR) posted through 31 March 2026, each charged to the approval whose write-off on that "
        "account it reverses; the Willamette hub's December 2025 lines loaded twice counted once.",
        "Denials for documents not received: counsellor denials in 2025 with that reason; the collaborative's reason "
        "list was renumbered on 1 July 2025 (to 30 June, code 2 was documents not received and code 1 income above "
        "the limit), which the denials that reconsiderations reopen show.",
        "Prepared by %s, %s, %s, from the collaborative's 2024 and 2025 extracts." % (
            K.PEOPLE["requester"], K.PEOPLE["requester_title"], K.FOUNDATION),
    ]
    nt.column_dimensions["A"].width = 120
    for i, t in enumerate(notes, 1):
        cell = nt.cell(i, 1, t)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            cell.font = Font(bold=True)
    wb.save(path)


def chart(path, fig):
    order = sorted(K.CODES, key=lambda x: fig["measure"][x])
    vals = [fig["measure"][x] for x in order]
    colors = ["#1D5C63" if fig["verdict"][x] == "Yes" else "#A9B1B8" for x in order]
    f, ax = plt.subplots(figsize=(8.4, 8.6), dpi=150)
    ax.barh(range(len(order)), vals, color=colors, height=0.66)
    ax.axvline(K.THRESHOLD, color="#B23A48", lw=1.4, ls="--")
    ax.text(K.THRESHOLD + 0.12, len(order) - 0.2, "6 percent: a coaching grant", color="#B23A48", fontsize=8.5,
            va="center")
    for i, v in enumerate(vals):
        ax.text(v + 0.12, i, "%.1f" % fig["measure_1dp"][order[i]], va="center", ha="left", fontsize=7.8,
                color="#1F2428")
    ax.set_yticks(range(len(order)))
    ax.set_yticklabels(["%s %s" % (x, K.NAME[x]) for x in order], fontsize=7.8)
    ax.set_xlim(0, max(vals) + 1.6)
    ax.set_xlabel("Eligible households' applications that left them owing more than half the balance, 2025 (%)",
                  fontsize=8.8)
    ax.set_title("%d of 31 hospitals reach six percent" % fig["call"]["count"], fontsize=12, loc="left",
                 color="#1F2428")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.tick_params(axis="x", labelsize=8.5)
    f.text(0.01, 0.008, "Counsellors' 2025 first-review decisions; eligibility on income over the twelve months before "
                        "applying,\nread through each region's calendar for each kind of work. %s, January 2027."
           % K.FOUNDATION, fontsize=7, color="#5D646A", linespacing=1.4)
    f.tight_layout(rect=(0, 0.035, 1, 1))
    f.savefig(path, metadata={"Software": None})
    plt.close(f)


def write_figures(path, fig):
    with open(path, "w") as fh:
        json.dump(fig, fh, indent=1, sort_keys=True)
        fh.write("\n")


def write(target, golden):
    os.makedirs(golden, exist_ok=True)
    res = A.run(target)
    fig = figures(res)
    workbook(os.path.join(golden, K.DELIVERABLES[0]), fig)
    chart(os.path.join(golden, K.DELIVERABLES[1]), fig)
    return fig, res
