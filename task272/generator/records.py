"""The pack's data files, written from the world and its operations: the three hubs' determination extracts (each in
its consolidation tool's own layout, codes and clock), H23's counsellor ID crosswalk, the member hospitals, the
assistance adjustments and the patient accounts they post to, and two files the analysis does not need (the patient
experience survey by payer and the hospitals' state charity care filings)."""
import csv
import datetime as dt
from zoneinfo import ZoneInfo

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import common as K

PACIFIC = ZoneInfo("America/Los_Angeles")
UTC = ZoneInfo("UTC")


def _w(path, header, rows, delimiter=","):
    with open(path, "w", newline="", encoding="utf-8") as f:
        wr = csv.writer(f, delimiter=delimiter, lineterminator="\r\n")
        wr.writerow(header)
        for r in rows:
            wr.writerow(r)
    return len(rows)


def reason_code(reason, day):
    return (K.REASON_V1 if day < K.REASON_SWITCH else K.REASON_V2)[reason]


def to_utc(local):
    return local.replace(tzinfo=PACIFIC).astimezone(UTC)


def _by_hub(rows, hub):
    sel = [r for r in rows if r["hub"] == hub]
    sel.sort(key=lambda r: (r["when"], r["code"], r["app"]))
    return sel


# ------------------------------------------------------------------------------------------- the hub extracts
def write_puget(path, rows):
    out = []
    for i, r in enumerate(_by_hub(rows, "puget")):
        typ = {"INITIAL": "INITIAL", "RECON": "RECONSIDERATION", "PRESUMPTIVE": "PRESUMPTIVE"}[r["type"]]
        out.append([r["code"], "D%08d" % (31000000 + i), r["app"], r["guarantor"], r["user"],
                    r["when"].strftime("%Y-%m-%d %H:%M:%S"), typ, K.PUGET_OUTCOME[r["outcome"]],
                    r["tier"] if r["outcome"] == "approve" else "",
                    reason_code(r["reason"], r["when"].date()) if r["outcome"] == "deny" else ""])
    return _w(path, ["hospital_code", "determination_id", "application_id", "guarantor_id", "user_id", "determined_at",
                     "determination_type", "outcome", "discount_pct", "denial_reason"], out)


def write_willamette(path, rows):
    sel = _by_hub(rows, "willamette")
    seq = {}
    out = []
    for r in sel:
        seq[r["code"]] = seq.get(r["code"], 0) + 1
        out.append([r["code"], "%d%06d" % (int(r["code"][1:]), seq[r["code"]]), r["app"], r["guarantor"], r["user"],
                    r["when"].strftime("%m/%d/%Y"), r["when"].strftime("%H:%M:%S"),
                    "Y" if r["type"] == "RECON" else "N", K.WILLAMETTE_OUTCOME[r["outcome"]],
                    r["tier"] if r["outcome"] == "approve" else "",
                    reason_code(r["reason"], r["when"].date()) if r["outcome"] == "deny" else ""])
    out.sort(key=lambda x: (x[0], int(x[1])))
    return _w(path, ["HOSP_ID", "DET_NO", "APPL_NO", "GUAR_NO", "USER_ID", "DET_DATE", "DET_TIME", "RECON", "RESULT",
                     "ADJ_PCT", "DENY_RSN"], out)


def write_columbia(path, rows):
    out = []
    for i, r in enumerate(_by_hub(rows, "columbia")):
        if r["outcome"] == "approve":
            dec = K.COLUMBIA_OUTCOME[("approve", r["tier"])]
        else:
            dec = K.COLUMBIA_OUTCOME[r["outcome"]]
        kind = {"INITIAL": "new", "RECON": "recheck", "PRESUMPTIVE": "screen"}[r["type"]]
        out.append([r["code"], r["app"], 5400000 + i, r["guarantor"], r["user"],
                    to_utc(r["when"]).strftime("%Y-%m-%dT%H:%M:%SZ"), kind, dec,
                    r["tier"] if r["outcome"] == "approve" else "",
                    reason_code(r["reason"], r["when"].date()) if r["outcome"] == "deny" else ""])
    return _w(path, ["facility", "app_ref", "seq", "guarantor", "staff_key", "ts_utc", "kind", "decision", "pct",
                     "reason"], out)


# ------------------------------------------------------------------------------------------- workbooks
def _xl_header(ws, row, cols, widths):
    for j, (c, wdt) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F4F6F")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = wdt


def write_crosswalk(path, crosswalk):
    wb = Workbook()
    ws = wb.active
    ws.title = "H23 counsellor IDs"
    ws["A1"] = "Lundgren Regional Hospital (H23): counsellor user IDs before and after the billing conversion"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("H23's billing moved from the Columbia hub's system to the Puget hub's on 2 June 2025. Each counsellor "
                "kept working under a new user ID; the Columbia staff key ends on 30 May 2025.")
    _xl_header(ws, 4, ["Hospital", "Columbia staff key (to 30 May 2025)", "Puget user ID (from 2 June 2025)",
                       "Effective"], [14, 30, 30, 14])
    for i, x in enumerate(crosswalk, 5):
        ws.cell(row=i, column=1, value=x["code"])
        ws.cell(row=i, column=2, value=x["old"])
        ws.cell(row=i, column=3, value=x["new"])
        c = ws.cell(row=i, column=4, value=x["effective"])
        c.number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A5"
    wb.save(path)
    return len(crosswalk)


def write_hospitals(path, hs, seed):
    r = np.random.default_rng([seed, 41])
    wb = Workbook()
    ws = wb.active
    ws.title = "Members 2025"
    ws["A1"] = "%s: member hospitals, 2025" % K.COLLAB
    ws["A1"].font = Font(bold=True, size=12)
    cols = ["Code", "Hospital", "City", "State", "Hub", "Licensed beds", "Free care up to (% of FPL)",
            "Discounts up to (% of FPL)", "Joined the collaborative"]
    _xl_header(ws, 3, cols, [8, 36, 16, 7, 34, 10, 12, 12, 12])
    for i, (code, name, city, state, hub) in enumerate(K.HOSPITALS, 4):
        hub_txt = {"puget": "Puget", "willamette": "Willamette", "columbia": "Columbia"}[hub]
        if code == K.MIGRATED:
            hub_txt = "Columbia to 30 May 2025; Puget from 2 June 2025"
        beds = int(r.integers(25, 120)) if state == "OR" and r.random() < 0.5 else int(r.integers(60, 480))
        free = 300 if state == "WA" and beds >= 200 else 200
        disc = 400 if state == "WA" else int(r.choice([300, 400]))
        joined = dt.date(int(r.choice([2023, 2023, 2024])), int(r.choice([1, 4, 7, 10])), 1)
        for j, v in enumerate([code, name, city, state, hub_txt, beds, free, disc, joined], 1):
            c = ws.cell(row=i, column=j, value=v)
            if j == 9:
                c.number_format = "yyyy-mm"
    ws.freeze_panes = "C4"
    ws.auto_filter.ref = "A3:I%d" % (3 + len(K.HOSPITALS))
    wb.save(path)
    return len(K.HOSPITALS)


# ------------------------------------------------------------------------------------------- adjustments and accounts
def write_ledger(path, ledger):
    out = []
    for t in ledger:
        row = [t["id"], t["code"], t["app"], t["account"], t["posted"].isoformat(), t["kind"], "%.2f" % t["amount"]]
        out.append(row)
        # the Willamette hub's December 2025 postings were loaded twice
        if K.HUB.get(t["code"]) == "willamette" and t["posted"].year == 2025 and t["posted"].month == 12:
            out.append(list(row))
    return _w(path, ["adjustment_id", "hospital_code", "application_id", "account_number", "posted_date",
                     "adjustment_code", "amount"], out)


def write_accounts(path, accounts, ledger):
    keep = sorted({t["account"] for t in ledger})
    out = []
    for a in keep:
        x = accounts[a]
        out.append([a, x["code"], x["guarantor"], x["service"].isoformat(), "%.2f" % x["balance"]])
    out.sort(key=lambda z: (z[1], z[0]))
    return _w(path, ["account_number", "hospital_code", "guarantor_id", "service_date", "patient_balance"], out)


# ------------------------------------------------------------------------------------------- two files not needed
def write_survey(path, seed):
    r = np.random.default_rng([seed, 43])
    out = []
    for code in K.CODES:
        base = r.uniform(62, 80)
        for payer, shift, n in (("Medicaid", -r.uniform(1, 9), r.integers(60, 400)),
                                ("Medicare", r.uniform(-2, 4), r.integers(150, 900)),
                                ("Commercial", r.uniform(-1, 5), r.integers(150, 1100)),
                                ("Self-pay or uninsured", -r.uniform(2, 12), r.integers(20, 160))):
            out.append([code, payer, int(n), "%.1f" % (base + shift)])
    return _w(path, ["hospital_code", "payer_group", "surveys_returned", "would_recommend_top_box_pct"], out)


def write_state(path, ledger, seed):
    """Each hospital's charity care as filed with its state for its own fiscal year 2025, at full charges."""
    r = np.random.default_rng([seed, 47])
    wb = Workbook()
    ws = wb.active
    ws.title = "FY2025 filings"
    ws["A1"] = "Hospital charity care as filed with the states, fiscal year 2025 (at full charges)"
    ws["A1"].font = Font(bold=True, size=12)
    _xl_header(ws, 3, ["Code", "Hospital", "State", "Fiscal year ends", "Charity care at charges ($)",
                       "Patients granted charity care"], [8, 36, 7, 14, 18, 14])
    for i, (code, name, city, state, hub) in enumerate(K.HOSPITALS, 4):
        fye = dt.date(2025, 6, 30) if (state == "OR" or r.random() < 0.6) else dt.date(2025, 12, 31)
        dollars = float(np.round(r.uniform(4.0, 30.0) * 1e6, -2))
        for j, v in enumerate([code, name, state, fye, dollars, int(r.integers(1500, 12000))], 1):
            c = ws.cell(row=i, column=j, value=v)
            if j == 4:
                c.number_format = "yyyy-mm-dd"
            if j == 5:
                c.number_format = "#,##0"
    ws.freeze_panes = "C4"
    wb.save(path)
    return len(K.HOSPITALS)
