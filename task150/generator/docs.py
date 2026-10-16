"""The governing document and the operating files around the records: Board Policy FEA-R2 (the estimand's filing), the
facilities capacity report, the student-system notes, the births file notes, the export manifest, and the two files
the solution never reads (the 2025 facility condition assessment and the October meal eligibility counts)."""
from __future__ import annotations

import datetime as dt
import os

import openpyxl
from openpyxl.styles import Alignment, Font
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

import common as K
import world as WD

F_POLICY = "BP_FEA-R2_elementary_classroom_placement.pdf"
F_CAPACITY = "facilities_capacity_report_2026-09.xlsx"
F_SIS_NOTES = "sis_export_notes.txt"
F_BIRTH_NOTES = "births_file_notes.txt"
F_MANIFEST = "export_manifest_2026-10-09.csv"
F_FCA = "facility_condition_assessment_2025.xlsx"
F_MEALS = "meal_eligibility_2026-10.csv"

BODY = ParagraphStyle("body", fontName="Times-Roman", fontSize=10.6, leading=13.6, alignment=TA_LEFT, spaceAfter=5)
HEAD = ParagraphStyle("head", fontName="Times-Bold", fontSize=10.8, leading=14, spaceBefore=7, spaceAfter=3,
                      keepWithNext=1)
TITLE = ParagraphStyle("title", fontName="Times-Bold", fontSize=13.5, leading=17, spaceAfter=3)
SMALL = ParagraphStyle("small", fontName="Times-Roman", fontSize=9, leading=11.2, textColor=colors.HexColor("#222222"))


def _pdf(path, story, title, subject, stamp):
    def on_page(c, d):
        c.setFont("Times-Roman", 8.5)
        c.drawString(0.9 * inch, 0.55 * inch, "Sorrel Creek RE-6 Board Policy Manual, Section F: Facilities Development")
        c.drawRightString(7.6 * inch, 0.55 * inch, "FEA-R2, page %d" % d.page)

    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.9 * inch, rightMargin=0.9 * inch, topMargin=0.8 * inch,
                            bottomMargin=0.85 * inch, title=title, author=K.DISTRICT, subject=subject,
                            creator=K.DISTRICT, producer=K.DISTRICT, invariant=1)
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    K.fix_pdf(path, stamp, "Board policy manual export")


def write_policy(out):
    s = []
    s.append(Paragraph("SORREL CREEK SCHOOL DISTRICT RE-6, BRANDT COUNTY, COLORADO", SMALL))
    s.append(Paragraph("Board of Education", SMALL))
    s.append(Spacer(1, 8))
    s.append(Paragraph("Policy FEA-R2: Elementary Classroom Placement (2025 Bond Program)", TITLE))
    s.append(Paragraph("Adopted March 12, 2026. Related policies: FB (Facilities Planning), JC (School Attendance "
                       "Areas).", SMALL))
    s.append(Spacer(1, 6))
    secs = [
        ("1. Program", [
            "This regulation governs where the 24 permanent classrooms funded by the Elementary Capacity program of the "
            "bond approved by the district's voters on November 4, 2025 are built. The classrooms are placed among the "
            "district's seven attendance-area elementary schools: Fifth Street, Lincoln, Ortega, Bluestem, Prairie Rose, "
            "Kessler and North Fork. They are built during the 2027-28 school year and serve the five school years "
            "2028-29 to 2032-33. Other bond programs (Facility Renewal, Safety and Security, Middle School Addition) are "
            "governed by their own regulations."]),
        ("2. Definitions", [
            "(a) K-5 enrollment of a school on a count day is the number of pupils in kindergarten to grade 5 whose "
            "primary enrollment is at that school on that day, each pupil counted once. Children enrolled in preschool "
            "are not counted. The count day is October 1, or the school day nearest October 1 when it falls on a "
            "weekend.",
            "(b) Projected enrollment of a school for a school year is its K-5 enrollment expected on that year's count "
            "day under the attendance areas in effect on October 1, 2026, including pupils who are not yet enrolled or "
            "not yet born, projected from the district's records through the October 1, 2026 count.",
            "(c) Five-year average projected enrollment is the mean of a school's projected enrollment for the school "
            "years 2028-29 to 2032-33.",
            "(d) Rated capacity is the K-5 rated capacity of the school's permanent classrooms in the facilities "
            "capacity report current on October 1, 2026. Portable classrooms are not counted.",
            "(e) Classroom need of a school is its five-year average projected enrollment less its rated capacity, "
            "divided by 24 and rounded up to a whole classroom. A school whose five-year average projected enrollment "
            "does not exceed its rated capacity has no classroom need."]),
        ("3. Placement", [
            "Schools with a classroom need are taken in descending order of the amount by which their five-year "
            "average projected enrollment exceeds their rated capacity. Each in turn is placed its classroom need or "
            "eight classrooms, whichever is fewer, or the classrooms that remain when fewer remain, until the 24 "
            "classrooms are placed. Two schools with the same amount are taken in order of their October 1, 2026 K-5 "
            "enrollment, the larger first. Classrooms that remain after every classroom need is met are held in "
            "reserve for the Board to place."]),
        ("4. Preparation and adoption", [
            "The Planning and Enrollment Office prepares the projection and the placement schedule from the district's "
            "student records and the county's birth records. The Superintendent presents the schedule to the Board for "
            "adoption at its regular November 2026 meeting. The schedule states, for each school, its classroom need "
            "and the classrooms placed."]),
        ("5. Changes after adoption", [
            "Attendance-area changes and enrollment counts after October 1, 2026 do not change an adopted schedule. "
            "Classrooms placed at a school are not moved to another school except by action of the Board."]),
        ("6. Review", [
            "The Board reviews this regulation when the Elementary Capacity program closes."]),
    ]
    for h, paras in secs:
        s.append(Paragraph(h, HEAD))
        for p in paras:
            s.append(Paragraph(p, BODY))
    s.append(Spacer(1, 10))
    s.append(Paragraph("History: adopted March 12, 2026 (Resolution 2026-07).", SMALL))
    _pdf(os.path.join(out, F_POLICY), s, "Policy FEA-R2 Elementary Classroom Placement (2025 Bond Program)",
         "Board Policy FEA-R2", dt.datetime(2026, 3, 13, 10, 4, 51))


# ---------------------------------------------------------------------------------------------------- capacity report
OPENED = {"P": 1952, "C": 1961, "A": 1988, "B": 1996, "D": 1999, "E": 1974, "F": 1983}
AREA_SQFT = {"P": 61850, "C": 54420, "A": 66310, "B": 68940, "D": 63275, "E": 74880, "F": 55190}


def capacity_rows(caps, plans, portables_used):
    """One row per school: permanent classrooms, other assigned rooms, K-5 teaching stations, rated capacity,
    portables on site (the ones in use plus any spare)."""
    r = K.rng("portables")
    rows = []
    for z in K.LIST_ORDER:
        pl = plans[z]
        on_site = portables_used[z] + int(r.integers(0, 3))     # spares, and portables used for other programs
        note = ""
        if z == "D":
            note = "Southeast Elementary until 06/30/2020"
        if z == "E":
            note = "Gym and kitchen remodel 2024; no change to classrooms"
        if z == "P":
            note = "Two portables relocated from Lincoln, summer 2023"
        rows.append([K.SCHOOLS[z]["state"], K.SCHOOLS[z]["name"], OPENED[z], AREA_SQFT[z], pl["permanent"],
                     pl["permanent"] - pl["stations"], pl["stations"], int(caps[z]), on_site, note])
    return rows


def write_capacity(out, caps, plans, portables_used):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Elementary 2026-27"
    ws.append(["%s, Facilities and Operations" % K.DISTRICT])
    ws.append(["Elementary school capacity, 2026-27 (rooms as assigned for the 2026-27 school year)"])
    ws.append(["Prepared by %s, Director of Facilities, 09/25/2026" % K.PEOPLE["facilities"]])
    ws.append([])
    hdr = ["State code", "School", "Opened", "Building area (sq ft)", "Permanent classrooms", "Other assigned rooms",
           "K-5 teaching stations", "Rated capacity (K-5)", "Portable classrooms on site", "Notes"]
    ws.append(hdr)
    rows = capacity_rows(caps, plans, portables_used)
    for row in rows:
        ws.append(row)
    n = 5 + len(rows)
    ws.append([])
    ws.append(["Rated capacity is the number of K-5 pupils the school's permanent K-5 teaching stations seat at the "
               "district's room loading (Facilities Master Plan 2019, section 4), assessed room by room."])
    ws.append(["Other assigned rooms: preschool, special education, intervention, art, music and STEM rooms as "
               "assigned for 2026-27. Portable classrooms are not included in rated capacity; some house before- and "
               "after-school programs or stand vacant."])
    ws["A1"].font = Font(bold=True, size=12)
    for c in ws[5]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
    ws.row_dimensions[5].height = 30
    for col, wd in zip("ABCDEFGHIJ", (10, 25, 8, 12, 12, 11, 11, 11, 12, 44)):
        ws.column_dimensions[col].width = wd
    for row in ws.iter_rows(min_row=6, max_row=n):
        row[3].number_format = "#,##0"
    ws.freeze_panes = "C6"
    K.save_workbook(wb, os.path.join(out, F_CAPACITY), dt.datetime(2026, 9, 25, 14, 18, 0), K.PEOPLE["facilities"],
                    "Elementary school capacity 2026-27")
    return rows


# ---------------------------------------------------------------------------------------------------- notes
def write_sis_notes(out, count_days):
    ds = [d.strftime("%m/%d/%Y") for d in count_days]
    days = "\n  " + "\n  ".join(", ".join(ds[k:k + 7]) for k in range(0, len(ds), 7))
    txt = """SORREL CREEK SCHOOL DISTRICT RE-6, STUDENT INFORMATION SYSTEMS
Notes on the student-system files pulled for the Planning and Enrollment Office, October 2026
Mariah Henry, Student Information Systems Manager, 10/08/2026

FILES
  october_membership_2019-2026.csv     October membership snapshots, current student system, 2019-20 to 2026-27
  membership_oct2013-oct2018.txt       October membership snapshots, legacy student system, 2013-14 to 2018-19
  student_number_crosswalk_2019.csv    legacy student numbers and the state IDs they converted to in 2019
  school_codes.csv                     elementary school codes in both systems
  sections_fall2026.csv                elementary sections scheduled for 2026-27
  section_enrollment_fall2026.csv      pupils' lines in those sections

MEMBERSHIP SNAPSHOTS
A snapshot is taken at the close of each October count day. It holds one row for each pupil whose primary
enrollment on that day is in kindergarten to grade 5 at one of the seven elementary schools. Children enrolled
in preschool are not in the snapshot. The count day is October 1, or the nearest school day when October 1
falls on a weekend. Count days:%s

Current system (comma-separated, ISO dates), 2019-20 on:
  school_year    school year of the count, e.g. 2026-27
  count_date     the count day
  school_code    state school code (school_codes.csv)
  school_name    the school's name on the count day
  student_id     state-assigned student ID (SASID), 10 digits
  grade_level    00 kindergarten, 01 to 05 grades 1 to 5
  birth_date
  gender         F, M; X available from 2021-22
  entry_date     the date the pupil's enrollment at this school began
  entry_type     01 entered the district in kindergarten; 02 from another Colorado district; 03 from another
                 state; 04 from a nonpublic school or home school; 05 from outside the United States;
                 06 from another school in the district
  ell            Y when the pupil is an English learner on the count day
  iep            Y when the pupil has an IEP on the count day

Legacy system (pipe-delimited, US dates), 2013-14 to 2018-19:
  SCHL_YR, CNT_DT   school year and count day
  SCH_CD            legacy school code (school_codes.csv)
  SCH_NAME          school name as the legacy system abbreviated it
  STU_NBR           legacy student number, kept while the pupil stayed in the district
  SASID             state ID; blank when the state's ID system had not returned one by the count day
  GRADE             KG, 01 to 05
  BIRTH_DT, SEX
  ENTRY_DT          the date the pupil's enrollment at this school began
  ENTRY_CD          E1 new to the district; E2 from another school in the district
  ELL, SPED         Y or N on the count day

THE 2019 CONVERSION
The district moved to the current student system in summer 2019. Legacy student numbers were converted to state
IDs from the 2018-19 roll: the crosswalk lists every pupil on that roll with the state ID the number converted to,
including pupils whose SASID was blank on the October 2018 snapshot. Pupils who left the district before 2018-19
were not converted. School codes changed from the legacy codes to the state codes, and entry codes were converted
to the current list for every enrollment carried into the new system.

SOUTHEAST ELEMENTARY
Southeast Elementary was renamed Kessler Elementary on 07/01/2020. Building, attendance area and codes did not
change (legacy 105, state 6175).

SECTIONS, FALL 2026
sections_fall2026.csv holds one row per section scheduled at the elementary schools for 2026-27: homeroom courses
HR-K (kindergarten) to HR-5, and the intervention courses INT-RDG, INT-MTH and INT-ELD. Each row gives the teacher
of record, the room and the session (FULL, AM, PM) the section meets in, and the section's start and end dates.
section_enrollment_fall2026.csv holds one row per pupil line in a section; student_id is the state ID, and a blank
end_date is a line still open at the pull (10/08/2026).

CLASS COUNTS
For class-size reporting the district counts classes, not sections. A homeroom class is one room in one session
(FULL, AM or PM), whatever the number of homeroom section rows or teachers of record in it, and it is counted in the
grade or grades of the pupils enrolled in it. A class meets on a day when one of its sections is open that day with
at least one pupil enrolled.
""" % days
    with open(os.path.join(out, F_SIS_NOTES), "w", encoding="utf-8", newline="\n") as f:
        f.write(txt)


def write_birth_notes(out):
    txt = """BRANDT COUNTY PUBLIC HEALTH, VITAL RECORDS
Resident live births by Sorrel Creek RE-6 attendance area, 2008 to 2025
Prepared for the RE-6 Planning and Enrollment Office under the 2021 data sharing agreement
Brian Owen, Registrar, 09/21/2026

What is counted. Live births to mothers resident in Brandt County, by year of birth and the mother's residence
on the birth certificate. Residence addresses are geocoded to the RE-6 attendance-area boundaries supplied by
Joseph Hanson in the district's GIS office (boundaries in effect 08/01/2026), and every year is coded to those
boundaries. Births to county residents outside the district are in the row OUTSIDE RE-6. Addresses that could
not be matched to a location are in the row NOT GEOCODED.

Status. 2008 to 2024 are final. 2025 is provisional: records registered by 03/31/2026. Registrations after that
date have added 1 to 2 percent to a year's final count in recent years.

File. brandt_county_births_by_attendance_area_2008-2025.csv: birth_year, attendance_area, births, status. Counts
are not suppressed in this file, which is provided under the agreement and not for publication.
"""
    with open(os.path.join(out, F_BIRTH_NOTES), "w", encoding="utf-8", newline="\n") as f:
        f.write(txt)


MANIFEST = [
    # file, source system, owner, contents, pulled
    ("october_membership_2019-2026.csv", "Student information system (current)", "Mariah Henry",
     "October membership snapshots 2019-20 to 2026-27, K-5", "10/08/2026"),
    ("membership_oct2013-oct2018.txt", "Student information system (legacy archive)", "Mariah Henry",
     "October membership snapshots 2013-14 to 2018-19, K-5", "10/08/2026"),
    ("student_number_crosswalk_2019.csv", "Student information system (2019 conversion)", "Mariah Henry",
     "Legacy student numbers on the 2018-19 roll and their state IDs", "10/08/2026"),
    ("school_codes.csv", "Student information system", "Mariah Henry",
     "Elementary school codes, legacy and state", "10/08/2026"),
    ("sections_fall2026.csv", "Student information system", "Mariah Henry",
     "Elementary sections scheduled for 2026-27", "10/08/2026"),
    ("section_enrollment_fall2026.csv", "Student information system", "Mariah Henry",
     "Pupil lines in the fall 2026 elementary sections", "10/08/2026"),
    ("sis_export_notes.txt", "Student Information Systems", "Mariah Henry",
     "Notes on the student-system files", "10/08/2026"),
    ("brandt_county_births_by_attendance_area_2008-2025.csv", "Brandt County Public Health, Vital Records",
     "Brian Owen", "Resident live births by RE-6 attendance area, 2008 to 2025", "09/21/2026"),
    ("births_file_notes.txt", "Brandt County Public Health, Vital Records", "Brian Owen",
     "Notes on the births file", "09/21/2026"),
    ("facilities_capacity_report_2026-09.xlsx", "Facilities and Operations", "Troy Medina",
     "Elementary school capacity, 2026-27", "09/25/2026"),
    ("facility_condition_assessment_2025.xlsx", "Facilities and Operations", "Troy Medina",
     "Facility condition assessment of the elementary schools, spring 2025", "10/06/2026"),
    ("BP_FEA-R2_elementary_classroom_placement.pdf", "Board policy manual", "Board of Education",
     "Policy FEA-R2, Elementary Classroom Placement (2025 Bond Program)", "10/06/2026"),
    ("preschool_enrollments_2026-27.csv", "Preschool registration platform", "Early Learning Office",
     "Preschool enrollments at the elementary sites, 2026-27", "10/07/2026"),
    ("ecse_roster_2026-10-01.xlsx", "Special Education Department", "Special Education Department",
     "Early childhood special education roster, active IEPs on 10/01/2026", "10/02/2026"),
    ("meal_eligibility_2026-10.csv", "Nutrition Services", "Nutrition Services",
     "Meal benefit eligibility of K-5 pupils on the October 1, 2026 count", "10/07/2026"),
]


def write_manifest(out, rows_by_file):
    rows = []
    for (f, sysname, owner, contents, pulled) in MANIFEST:
        rows.append([f, sysname, owner, contents, pulled, rows_by_file.get(f, "")])
    return K.write_csv(os.path.join(out, F_MANIFEST),
                       ["file_name", "source", "owner", "contents", "received", "records"], rows)


# ---------------------------------------------------------------------------------------------------- distractors
def write_fca(out):
    """Facilities' 2025 condition assessment of the elementary buildings: condition scores and deferred maintenance by
    building system (a Facility Renewal program input)."""
    r = K.rng("fca")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Elementary FCA 2025"
    ws.append(["%s, Facilities and Operations" % K.DISTRICT])
    ws.append(["Facility condition assessment, elementary schools, spring 2025 (Facility Renewal program planning)"])
    ws.append([])
    hdr = ["State code", "School", "Assessed", "Current replacement value ($)", "Deferred maintenance ($)", "FCI",
           "Roofing", "HVAC", "Electrical", "Plumbing", "Interiors", "Site and paving", "Accessibility"]
    ws.append(hdr)
    rating = ["Good", "Fair", "Fair", "Poor"]
    for z in K.LIST_ORDER:
        age = 2025 - OPENED[z]
        crv = int(round(AREA_SQFT[z] * float(r.uniform(388, 421)), -3))
        fci = min(0.46, 0.04 + age * float(r.uniform(0.0042, 0.0061)))
        dm = int(round(crv * fci, -2))
        sc = []
        for k in range(7):
            p = min(3, max(0, int(round(age / 22 + float(r.normal(0, 0.8))))))
            sc.append(rating[p])
        assessed = dt.date(2025, 3, 3) + dt.timedelta(days=int(r.integers(0, 70)))
        ws.append([K.SCHOOLS[z]["state"], K.SCHOOLS[z]["name"], assessed, crv, dm, round(dm / crv, 3)] + sc)
    for row in ws.iter_rows(min_row=5):
        row[2].number_format = "mm/dd/yyyy"
        row[3].number_format = "#,##0"
        row[4].number_format = "#,##0"
        row[5].number_format = "0.000"
    ws.append([])
    ws.append(["FCI is deferred maintenance divided by current replacement value. Ratings by building system: Good, "
               "Fair, Poor."])
    ws["A1"].font = Font(bold=True, size=12)
    for c in ws[4]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
    for col, wd in zip("ABCDEFGHIJKLM", (10, 25, 11, 16, 15, 7, 9, 9, 10, 9, 9, 10, 12)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "C5"
    K.save_workbook(wb, os.path.join(out, F_FCA), dt.datetime(2025, 6, 13, 11, 2, 0), K.PEOPLE["facilities"],
                    "Elementary facility condition assessment 2025")


FRL = {"A": (0.47, 0.11), "B": (0.15, 0.06), "P": (0.53, 0.12), "C": (0.22, 0.08), "D": (0.19, 0.07),
       "E": (0.26, 0.09), "F": (0.11, 0.05)}


def write_meals(W, out):
    """Nutrition Services' October eligibility counts for the K-5 pupils on the October 1, 2026 count."""
    r = K.rng("meals")
    rows = []
    for z in K.LIST_ORDER:
        n = int(W.counts[2026][WD.Z.index(z)].sum())
        f, rd = FRL[z]
        free = int(r.binomial(n, f))
        red = int(r.binomial(n - free, rd / (1 - f)))
        direct = int(r.binomial(free, 0.58))
        rows.append([K.SCHOOLS[z]["state"], K.SCHOOLS[z]["name"], n, free, red, n - free - red, direct,
                     "%.1f" % (100.0 * (free + red) / n)])
    K.write_csv(os.path.join(out, F_MEALS),
                ["school_code", "school_name", "k5_pupils_oct1", "free", "reduced", "paid", "direct_certified",
                 "pct_free_reduced"], rows)
    return rows
