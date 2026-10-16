"""task150 checks, run on every build: the input gates; the three-rules block (R1 the field's raw material, R2 every
mess item neutral, R3a the analyst test on the clean truth and the pack, R3b the decisive step alone, R3c no plant,
R3s the strict rule 3 test on every shipped sentence that bears on how the figure is computed); the bins and the
one-sigma test; every construction of the golden; the out-of-sample rejection of the default and the rivals; every
rung by name; every semantic fork row (SF1 to SF16) and closure axis (C1 to C22); both asks with every stop, hazard and
the cracker's path; the separation line and the distractors; the realism sweep; the goldens and the prompt. Any
failure stops the build."""
from __future__ import annotations

import copy
import csv
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import tempfile
import zipfile

import numpy as np
import openpyxl

import asks as AS
import calc
import common as K
import deliver as DL
import docs as DOC
import records as REC
import world as WD

DASH = re.compile("[%s%s]" % (chr(0x2013), chr(0x2014)))
REPO = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
TASK = os.path.join(REPO, "task150") if os.path.isdir(os.path.join(REPO, "task150")) else K.TASK
CODE = {z: K.SCHOOLS[z]["state"] for z in K.ZONES}
DEFAULT_PLACED = {"6111": 3, "6116": 6, "6142": 0, "6158": 7, "6163": 8, "6175": 0, "6189": 0}
GOLDEN_PLACED = {"6111": 5, "6116": 6, "6142": 0, "6158": 4, "6163": 7, "6175": 2, "6189": 0}
GOLDEN_NEED = {"6111": 5, "6116": 6, "6142": 0, "6158": 4, "6163": 7, "6175": 3, "6189": 1}
MAIN_FILES = [REC.F_LEGACY, REC.F_MEMBERS, REC.F_XWALK, REC.F_CODES, REC.F_BIRTHS, DOC.F_CAPACITY, DOC.F_POLICY]
ASK_FILES = [AS.F_SECTIONS, AS.F_SECT_ENR, AS.F_PRESCHOOL, AS.F_ECSE]
ASOF = dt.date(2026, 10, 9)


class Checks:
    def __init__(self):
        self.res = []

    def __call__(self, label, ok, detail=""):
        self.res.append((label, bool(ok), detail))
        return ok

    def report(self):
        if os.environ.get("CHECKS_LABELS"):
            with open(os.environ["CHECKS_LABELS"], "w", encoding="utf-8") as fh:
                fh.writelines(("ok   " if ok else "FAIL ") + lab + "\n" for lab, ok, _ in self.res)
        bad = [r for r in self.res if not r[1]]
        for lab, ok, det in self.res:
            if not ok:
                print("FAIL %s %s" % (lab, str(det)[:900]))
        print("checks.py: %d assertions, %d failed" % (len(self.res), len(bad)))
        if bad:
            raise SystemExit(1)


# ---------------------------------------------------------------------------------------------------- readers
def pdf_text(path):
    out = subprocess.run(["pdftotext", "-enc", "UTF-8", path, "-"], capture_output=True, check=True).stdout
    return re.sub(r"\s+", " ", out.decode("utf-8"))


def xlsx_strings(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    bits = []
    for ws in wb.worksheets:
        bits.append(ws.title)
        for row in ws.iter_rows(values_only=True):
            bits.extend(str(v) for v in row if isinstance(v, str))
    return " ".join(bits)


def shipped_text(target):
    """Every sentence and header the pack ships."""
    out = {}
    for f in sorted(os.listdir(target)):
        p = os.path.join(target, f)
        if f.endswith(".pdf"):
            out[f] = pdf_text(p)
        elif f.endswith(".xlsx"):
            out[f] = xlsx_strings(p)
        elif f.endswith(".txt") and "membership" not in f:
            out[f] = open(p, encoding="utf-8").read()
        else:
            with open(p, encoding="utf-8", newline="") as fh:
                first = fh.readline()
            out[f] = first
            if f == DOC.F_MANIFEST:
                out[f] = open(p, encoding="utf-8").read()
    return out


def placed(D, a):
    return {c: int(a["alloc"][i]) for i, c in enumerate(D.codes)}


def needs(D, a):
    return {c: int(a["need"][i]) for i, c in enumerate(D.codes)}


def alloc(D, P):
    return calc.allocate(D, calc.avg5(D, P), D.cap)


# ---------------------------------------------------------------------------------------------------- the clean truth
def truth_data(W, Dp):
    """The world before any record is written: every pupil by his own key, the true births (no geocoding misses, no
    late registrations), the shipped rated capacities, in the calc's school order."""
    order = [WD.Z.index(z) for z in sorted(K.ZONES, key=lambda z: K.SCHOOLS[z]["state"])]
    inv = {zi: k for k, zi in enumerate(order)}
    D = calc.Data()
    D.codes, D.names, D.NS = Dp.codes, Dp.names, Dp.NS
    D.snap = {y: {pid: (inv[s], g) for (pid, s, g) in W.snap[y]} for y in K.YEARS}
    D.years = sorted(D.snap)
    D.E = {}
    for y in D.years:
        E = np.zeros((D.NS, calc.G))
        for (i, g) in D.snap[y].values():
            E[i, g] += 1
        D.E[y] = E
    D.births = {k: {y: W.births[WD.Z[zi]][y] for y in range(2004, 2026)} for k, zi in enumerate(order)}
    D.cap = Dp.cap
    D.rows = {y: len(D.snap[y]) for y in D.years}
    return D


def with_births(D, fn):
    D2 = copy.copy(D)
    D2.births = {i: {y: fn(i, y, b) for y, b in D.births[i].items()} for i in D.births}
    return D2


def kessler_post_rename(D, F):
    """The rename mishandled: Kessler's rates from the Octobers that carry its new name only (2020 on)."""
    prm = calc.gross_params(D, F, calc.ALLW, calc.ALLW)
    k = D.codes.index("6175")
    post = [y for y in calc.ALLW if y >= 2021]
    p2 = calc.gross_params(D, F, post, post)
    for key in ("outr", "kr"):
        prm[key] = prm[key].copy()
        prm[key][k] = p2[key][k]
    for key in ("mvr", "retr", "arrc"):
        prm[key] = prm[key].copy()
        prm[key][k] = p2[key][k]
    return calc.roll_gross(D, D.E[calc.LAST], prm)


# ---------------------------------------------------------------------------------------------------- the checks
def run(target, golden, W, R, meta_path):
    C = Checks()
    M, D, A, B = R["M"], R["D"], R["A"], R["B"]
    F = M["F"]
    g, d = M["golden"], M["default"]
    files = sorted(os.listdir(target))
    meta = json.load(open(meta_path, encoding="utf-8"))
    texts = shipped_text(target)

    # ---- input gates
    C("G1 ten or more files", len(files) >= 10, len(files))
    C("G2 three or more formats", len({f.rsplit(".", 1)[1] for f in files}) >= 3, files)
    big = max(sum(1 for _ in open(os.path.join(target, f), encoding="utf-8")) - 1
              for f in files if f.endswith(".csv") or f.endswith(".txt"))
    C("G3 a file of 25,000 rows or more", big >= 25000, big)
    dis = meta["distractor_files"]
    C("G4 two distractors named in metadata.json and shipped", len(dis) >= 2 and all(x in files for x in dis), dis)
    C("G4b no file name says distractor", not any("distract" in f.lower() for f in files))
    with open(os.path.join(target, DOC.F_MANIFEST), encoding="utf-8", newline="") as fh:
        man = [r["file_name"] for r in csv.DictReader(fh)]
    C("G5 the manifest lists every shipped file but itself (H9)", set(man) == set(files) - {DOC.F_MANIFEST},
      set(man) ^ (set(files) - {DOC.F_MANIFEST}))
    C("G6 every file carries source, date and licence in metadata.json",
      {x["path"] for x in meta["files"]} == set(files) and all(x["source"] and x["date"] and x["license"]
                                                              for x in meta["files"]))

    # ---- R1 the field's raw material
    C("R1a the decisive computation reads linked pupil rows: K-4 bases equal the membership counts",
      all(np.allclose(F[y]["base"], D.E[y - 1][:, :5].sum(axis=1)) for y in F))
    flowwords = re.compile(r"\b(rate|rates|ratio|ratios|projection|projected|flow|flows|arrivals?|exits?|survival|"
                           r"progression|transfers?)\b", re.I)
    hdrs = {f: texts[f] for f in files if f.endswith(".csv") and f != DOC.F_MANIFEST}
    C("R1b no shipped column holds a flow, rate, ratio or projection",
      not any(flowwords.search(h) for h in hdrs.values()), {f: h for f, h in hdrs.items() if flowwords.search(h)})
    C("R1c the membership exports hold pupil rows (one per pupil per count day), not counts",
      sum(D.rows.values()) > 70000 and all(len(D.snap[y]) == D.rows[y] for y in D.years))

    # ---- R2 every mess item neutral
    Dl = calc.load(target, drop_legacy=True)
    Fl = calc.flows(Dl)
    w = list(range(2020, 2027))
    a_l = alloc(Dl, calc.proj_gross(Dl, Fl, w, w))
    C("R2a two layouts: the 2019-26 export alone files the golden", placed(Dl, a_l) == GOLDEN_PLACED, placed(Dl, a_l))
    Ds = calc.load(target, link="sasid")
    Fs = calc.flows(Ds)
    a_s = alloc(Ds, calc.proj_gross(Ds, Fs))
    C("R2b blank SASIDs: linking by the SASID on the row (blanks unlinked) files the golden",
      placed(Ds, a_s) == GOLDEN_PLACED and needs(Ds, a_s) == GOLDEN_NEED, placed(Ds, a_s))
    blanks = sum(1 for r in csv.DictReader(open(os.path.join(target, REC.F_LEGACY), encoding="utf-8", newline=""),
                                           delimiter="|") if not r["SASID"])
    C("R2b' the legacy export carries blank SASIDs (a few per cent of first-year rows)", 100 < blanks < 600, blanks)
    a_k = alloc(D, kessler_post_rename(D, F))
    C("R2c the rename: Kessler fitted on its post-2020 Octobers only files the golden",
      placed(D, a_k) == GOLDEN_PLACED and needs(D, a_k) == GOLDEN_NEED, placed(D, a_k))
    cd = {}
    with open(os.path.join(target, REC.F_MEMBERS), encoding="utf-8", newline="") as fh:
        for r in csv.DictReader(fh):
            cd[r["school_year"]] = r["count_date"]
    C("R2d count days off October 1: keying snapshots by the count date's year is the school year",
      all(int(v[:4]) == int(k[:4]) for k, v in cd.items()) and len(cd) == 8)
    Dg = with_births(D, lambda i, y, b: b * (1 + ng_share(target, y)))
    a_g = alloc(Dg, calc.proj_gross(Dg, F))
    C("R2e births not geocoded spread over the areas pro rata files the golden (SF12)",
      placed(Dg, a_g) == GOLDEN_PLACED and needs(Dg, a_g) == GOLDEN_NEED, placed(Dg, a_g))
    Dp25 = with_births(D, lambda i, y, b: b * 1.016 if y == 2025 else b)
    a_p = alloc(Dp25, calc.proj_gross(Dp25, F))
    C("R2f 2025 births completed for late registrations files the golden (SF11)",
      placed(Dp25, a_p) == GOLDEN_PLACED and needs(Dp25, a_p) == GOLDEN_NEED, placed(Dp25, a_p))

    # ---- R3 the analyst test
    T = truth_data(W, D)
    FT = calc.flows(T)
    t_def = alloc(T, calc.proj_gpr(T, calc.W5))
    t_gold = alloc(T, calc.proj_gross(T, FT))
    C("R3a the default path on the clean truth files the named wrong answer", placed(T, t_def) == DEFAULT_PLACED,
      placed(T, t_def))
    C("R3a' the default path on the pack files the same named wrong answer", placed(D, d) == DEFAULT_PLACED,
      placed(D, d))
    C("R3a'' the golden on the clean truth", placed(T, t_gold) == GOLDEN_PLACED and needs(T, t_gold) == GOLDEN_NEED,
      (placed(T, t_gold), needs(T, t_gold)))
    t_dec = alloc(T, calc.proj_gross(T, FT, calc.W5, calc.W5))
    p_dec = M["alloc"]["last five"]
    C("R3b the default path with only the decisive step changed (flow by flow, same five-year window) files the golden",
      placed(T, t_dec) == GOLDEN_PLACED and placed(D, p_dec) == GOLDEN_PLACED, (placed(T, t_dec), placed(D, p_dec)))
    lower = ["three-year progression ratios", "eight-year progression ratios", "all-year progression ratios",
             "ratios trended over eight years", "cohort change differences, five years",
             "five-year ratios raked to a district total", "district ratios split by school share"]
    C("R3b' each lower rung's single correction leaves a wrong schedule",
      all(placed(D, M["alloc"][k]) != GOLDEN_PLACED for k in lower), [k for k in lower if placed(D, M["alloc"][k]) == GOLDEN_PLACED])
    # R3c no plant: no copied or backfilled block, no swapped series
    vecs = {}
    dup = []
    for y in D.years:
        for i in range(D.NS):
            v = tuple(D.E[y][i].astype(int))
            if v in vecs:
                dup.append((y, i, vecs[v]))
            vecs[v] = (y, i)
    C("R3c no grade vector repeats across schools and years (no copied block)", not dup, dup[:5])
    yr_tot = [tuple(int(D.E[y][i].sum()) for i in range(D.NS)) for y in D.years]
    rep_tot = sum(1 for k in range(len(yr_tot) - 1) for i in range(D.NS) if yr_tot[k][i] == yr_tot[k + 1][i])
    C("R3c' a school's total repeats in consecutive years at most once (a coincidence, not a block)", rep_tot <= 1,
      rep_tot)
    method = re.compile(r"\b(progression|survival|cohort|ratios?|rates?|arrivals?|flows?|trend\w*|extrapolat\w*|"
                        r"retention|transfers?|migration rates?|in-migration|out-migration)\b", re.I)
    hits = {f: sorted(set(m.group(0).lower() for m in method.finditer(t))) for f, t in texts.items()
            if not f.endswith(".csv") or f == DOC.F_MANIFEST}
    hits = {f: h for f, h in hits.items() if h}
    C("R3s no shipped sentence names a projection method, a ratio, a rate or how arrivals or moves behave",
      not hits, hits)
    pol = texts[DOC.F_POLICY]
    how = [s for s in re.split(r"(?<=\.)\s", pol) if re.search(r"project|prepar", s, re.I)]
    C("R3s' the policy's HOW sentences name records, never a method",
      len(how) >= 2 and all(not method.search(s) for s in how), how)
    gpr_family = ["five-year progression ratios (default)", "three-year progression ratios",
                  "latest-year progression ratios", "eight-year progression ratios", "all-year progression ratios",
                  "five-year mean of ratios", "ratios trended over eight years", "ratios trended over five years",
                  "cohort change differences, five years", "cohort change differences, all years",
                  "district ratios split by school share", "five-year ratios raked to a district total"]
    C("R3s'' the field's standard projection executed with every filed rule never files the golden",
      all(placed(D, M["alloc"][k]) != GOLDEN_PLACED for k in gpr_family))

    # ---- bins, constructions, one sigma
    C("B1 the golden schedule and needs", placed(D, g) == GOLDEN_PLACED and needs(D, g) == GOLDEN_NEED,
      (placed(D, g), needs(D, g)))
    cons = M["constructions"]
    C("B2 every construction of the golden files one schedule and one set of needs (C1)",
      all(placed(D, M["alloc"][k]) == GOLDEN_PLACED and needs(D, M["alloc"][k]) == GOLDEN_NEED for k in cons),
      [k for k in cons if placed(D, M["alloc"][k]) != GOLDEN_PLACED or needs(D, M["alloc"][k]) != GOLDEN_NEED])
    ex = np.array([M["alloc"][k]["exc"] for k in cons])
    lo_e = np.floor(g["exc"] / 24.0) * 24.0
    inside = all(((ex[:, i] > lo_e[i]) & (ex[:, i] <= lo_e[i] + 24)).all() for i in range(D.NS) if g["exc"][i] > 0)
    C("B3 every construction's excess sits inside the golden's bin at every school with a need", inside,
      np.round(ex, 1))
    sd = M["sd"]
    margin = {D.codes[i]: (round(float(min(g["exc"][i] - lo_e[i], lo_e[i] + 24 - g["exc"][i])), 1), round(float(sd[i]), 1))
              for i in range(D.NS) if g["exc"][i] > 0}
    C("B4 one sigma: every school's golden excess sits at least one standard error inside its bin",
      all(m >= s for m, s in margin.values()), margin)
    C("B4' Ortega's five-year average sits more than three standard errors below its rated capacity",
      g["exc"][D.codes.index("6142")] < -3 * sd[D.codes.index("6142")])
    C("B5 the golden's mid-bin distance is at least 7 pupils at every school with a need",
      all(m >= 7 for m, s in margin.values()), margin)

    # ---- the ladder and the out-of-sample rejection
    ladder = {
        "three-year progression ratios": {"6111": 4, "6116": 6, "6158": 6, "6163": 8, "6175": 0},
        "latest-year progression ratios": None, "ratios trended over eight years": None,
        "cohort change differences, five years": None, "five-year ratios raked to a district total": None,
        "moves in as counts": {"6111": 5, "6116": 6, "6158": 6, "6163": 7, "6175": 0},
        "arrivals as a share of the receiving class": {"6111": 4, "6116": 6, "6158": 4, "6163": 8, "6175": 2},
    }
    for k, want in ladder.items():
        got = placed(D, M["alloc"][k])
        C("L1 rung %s places another schedule%s" % (k, " (as the ladder states)" if want else ""),
          got != GOLDEN_PLACED and (want is None or all(got[c] == v for c, v in want.items())), got)
    C("L2 every rival places another schedule", all(placed(D, M["alloc"][k]) != GOLDEN_PLACED for k in M["rivals"]),
      [k for k in M["rivals"] if placed(D, M["alloc"][k]) == GOLDEN_PLACED])
    moved = sum(1 for c in GOLDEN_PLACED if GOLDEN_PLACED[c] != DEFAULT_PLACED[c])
    C("L3 the decisive step moves four of the seven placements", moved == 4, moved)
    bt = M["backtest"]
    ix = {c: i for i, c in enumerate(D.codes)}
    gl = bt["gross-flow"]["last"]
    C("L4 out of sample from 2019 the gross-flow projection misses every school's 2026 count by under 5.5 per cent",
      np.abs(gl).max() < 5.5, np.round(gl, 1))
    gp = bt["five-year progression ratios"]["last"]
    C("L4' five-year ratios come out over 7 per cent low at Fifth Street and over 7 per cent high at Prairie Rose",
      gp[ix["6111"]] < -7 and gp[ix["6163"]] > 7 and abs(gl[ix["6111"]]) < 2 and abs(gl[ix["6163"]]) < 2,
      (np.round(gp, 1), np.round(gl, 1)))
    er = bt["arrivals as a share"]["last"]
    C("L4'' arrivals as a share come out over 5 per cent low at Fifth Street and high at Prairie Rose",
      er[ix["6111"]] < -5 and er[ix["6163"]] > 5, np.round(er, 1))
    ic = bt["moves in as counts"]["mape"]
    gm = bt["gross-flow"]["mape"]
    C("L4''' moves in as counts miss Bluestem out of sample at more than twice the gross-flow projection's error",
      ic[ix["6158"]] > 2 * gm[ix["6158"]], (round(float(ic[ix["6158"]]), 2), round(float(gm[ix["6158"]]), 2)))
    cc = bt["cohort change differences"]["last"]
    C("L4'''' cohort change differences come out over 5 per cent low at Fifth Street", cc[ix["6111"]] < -5,
      np.round(cc, 1))
    fc = DL.facts(R)
    C("L5 Fifth Street's arrivals hold between 60 and 95 a year while its K-4 base falls by more than 10 per cent",
      min(fc["arr"]["6111"]) >= 60 and max(fc["arr"]["6111"]) <= 95 and fc["base"]["6111"][-1] < 0.9 * fc["base"]["6111"][0])
    C("L5' Prairie Rose's arrivals hold between 55 and 90 a year while its K-4 base rises by more than 30 per cent",
      min(fc["arr"]["6163"]) >= 55 and max(fc["arr"]["6163"]) <= 90 and fc["base"]["6163"][-1] > 1.3 * fc["base"]["6163"][0])
    r_ab = np.corrcoef(fc["a2b"], [F[y]["base"][ix["6142"]] for y in calc.ALLW])[0, 1]
    C("L5'' Ortega's moves to Bluestem rise and fall with Ortega's base (correlation over 0.6) at 9 to 12 per 100",
      r_ab > 0.6 and min(fc["a2b_rate"]) > 9 and max(fc["a2b_rate"]) < 12, (round(float(r_ab), 2), fc["a2b_rate"]))
    r_pa = np.corrcoef(fc["arr"]["6111"], fc["base"]["6111"])[0, 1]
    C("L5''' Fifth Street's arrivals do not follow its base (correlation under 0.5)", r_pa < 0.5, round(float(r_pa), 2))

    # ---- the semantic fork table
    C("SF1 the 24 classrooms are the Elementary Capacity program's", "24 permanent classrooms" in pol and
      "Elementary Capacity program" in pol)
    names = [n for n in D.names]
    C("SF2 seven attendance-area schools in the code list, the capacity report and the policy",
      len(D.codes) == 7 and all(n.replace(" Elementary", "") in pol for n in names))
    k_new = [sum(M["P"][calc.GOLDEN][y][:, 0]) for y in calc.WIN]
    C("SF3 projected enrollment includes pupils not yet enrolled or born (the policy's words; new kindergartens in "
      "every window year)", "not yet enrolled or not yet born" in pol and min(k_new) > 500)
    C("SF4 the window is 2028-29 to 2032-33", "2028-29 to 2032-33" in pol and calc.WIN == list(range(2028, 2033)))
    C("SF5 the average is the mean of five counts", "mean of a school's projected enrollment" in pol)
    C("SF6 rated capacity counts permanent classrooms, portables not", "Portable classrooms are not counted" in pol)
    C("SF7 the classroom need is uncapped (Kessler's need 3, placed 2)", g["need"][ix["6175"]] == 3 and
      g["alloc"][ix["6175"]] == 2 and "whichever is fewer" in pol)
    C("SF8 the placement follows the order, need and cap until 24", int(g["alloc"].sum()) == 24 and
      max(g["alloc"]) <= 8)
    grades_m = {r["grade_level"] for r in csv.DictReader(open(os.path.join(target, REC.F_MEMBERS), encoding="utf-8"))}
    C("SF9 K-5 only: the membership holds grades 00 to 05 and no preschool", grades_m == {"00", "01", "02", "03", "04", "05"})
    C("SF10 the projection starts from the October 1, 2026 count",
      np.allclose(M["P"][calc.GOLDEN][2026], D.E[2026]) and "through the October 1, 2026 count" in pol)
    C("SF11 2025 births provisional or completed: same call", placed(Dp25, a_p) == GOLDEN_PLACED)
    C("SF12 births not geocoded kept out or spread: same call", placed(Dg, a_g) == GOLDEN_PLACED)
    C("SF13 the expected count: one projection, the policy's word", "expected" in pol)
    avg_r = np.round(g["avg"])
    need_r = [max(0, int(np.ceil((avg_r[i] - D.cap[i]) / 24))) if avg_r[i] > D.cap[i] else 0 for i in range(D.NS)]
    C("SF14 rounding the average first gives the same needs", need_r == [int(x) for x in g["need"]], need_r)
    ex_sorted = sorted([float(e) for e in g["exc"] if e > 0], reverse=True)
    C("SF15 no tie in the order (every gap over 5 pupils)",
      all(ex_sorted[k] - ex_sorted[k + 1] > 5 for k in range(len(ex_sorted) - 1)), ex_sorted)
    C("SF16 the needs total more than 24, so nothing is held in reserve", int(g["need"].sum()) > 24 and g["left"] == 0,
      int(g["need"].sum()))

    # ---- the closure axes
    C("C1 population: kindergartens of 2028 to 2032 come from births 2023 to 2027 (none enrolled yet)",
      all(M["P"][calc.GOLDEN][y][:, 0].sum() > 0 for y in calc.WIN))
    uniq = all(len(D.snap[y]) == D.rows[y] for y in D.years)
    C("C2 unit: each pupil once per count day in both layouts", uniq)
    days = {}
    for r in csv.DictReader(open(os.path.join(target, REC.F_LEGACY), encoding="utf-8", newline=""), delimiter="|"):
        days[int(r["SCHL_YR"][:4])] = dt.datetime.strptime(r["CNT_DT"], "%m/%d/%Y").date()
    for k2, v in cd.items():
        days[int(k2[:4])] = dt.date.fromisoformat(v)
    C("C3 attribution: every snapshot sits on its filed count day", all(days[y] == K.COUNT_DAY[y] for y in K.YEARS),
      days)
    late = []
    for f in files:
        if not f.endswith(".csv") or f == DOC.F_MANIFEST:
            continue
        for r in csv.DictReader(open(os.path.join(target, f), encoding="utf-8", newline="")):
            for k2, v in r.items():
                if v and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v) and dt.date.fromisoformat(v) > ASOF:
                    if not (f == AS.F_SECTIONS and k2 == "section_end"):
                        late.append((f, k2, v))
    C("C4 as-of: no record dated after the pull except the sections' planned end dates (H16 allow-list)", not late,
      late[:5])
    C("C5 version basis: one snapshot per school year, 14 in all", len(D.years) == 14 and D.years == K.YEARS)
    C("C6 denominator: rated capacity for all seven schools from the capacity report",
      D.cap is not None and len(D.cap) == 7 and all(D.cap > 500))
    C("C7 weighting: exits and moves pooled over grades or by grade, the same schedule",
      placed(D, M["alloc"]["exits and moves by grade"]) == GOLDEN_PLACED)
    C("C8 window length: all, ten, eight, five or three transitions, the same schedule", all(
        placed(D, M["alloc"][k]) == GOLDEN_PLACED for k in ("all transitions", "last ten", "last eight", "last five",
                                                            "last three")))
    on_count_day = []
    for r in csv.DictReader(open(os.path.join(target, REC.F_MEMBERS), encoding="utf-8", newline="")):
        if r["entry_date"] == r["count_date"]:
            on_count_day.append(r["student_id"])
    C("C9 boundary: no enrollment begins on a count day", not on_count_day, on_count_day[:5])
    C("C10 rounding path: the need from the unrounded average (SF14)", need_r == [int(x) for x in g["need"]])
    C("C11 tie-break never invoked (SF15)", len(set(np.round(g["exc"], 6))) == 7)
    C("C12 maturity: every count day complete, 2025 births provisional (SF11)",
      all(D.rows[y] > 5000 for y in D.years))
    C("C13 order of operations: need from excess, placement from need", all(
        (g["need"][i] == (int(np.ceil(g["exc"][i] / 24)) if g["exc"][i] > 0 else 0)) for i in range(D.NS)))
    rev = list(csv.DictReader(open(os.path.join(target, REC.F_MEMBERS), encoding="utf-8", newline="")))[::-1]
    E26 = np.zeros((7, 6))
    for r in rev:
        if r["school_year"] == "2026-27":
            E26[D.codes.index(r["school_code"]), int(r["grade_level"])] += 1
    C("C14 row order: the export read backwards gives the same October 2026 counts", np.allclose(E26, D.E[2026]))
    leg_keys = [(r["SCHL_YR"], r["STU_NBR"]) for r in csv.DictReader(
        open(os.path.join(target, REC.F_LEGACY), encoding="utf-8", newline=""), delimiter="|")]
    mem_keys = [(r["school_year"], r["student_id"]) for r in rev]
    C("C15 duplicates: one row per pupil per count day in both exports",
      len(set(leg_keys)) == len(leg_keys) and len(set(mem_keys)) == len(mem_keys))
    xw = {r["STU_NBR"]: r["STATE_ID"] for r in csv.DictReader(open(os.path.join(target, REC.F_XWALK), encoding="utf-8"))}
    l18 = [r for r in csv.DictReader(open(os.path.join(target, REC.F_LEGACY), encoding="utf-8", newline=""),
                                     delimiter="|") if r["SCHL_YR"] == "2018-19"]
    C("C16 identity: the crosswalk covers the 2018-19 roll and agrees with every SASID on it",
      all(r["STU_NBR"] in xw for r in l18) and all(xw[r["STU_NBR"]] == r["SASID"] for r in l18 if r["SASID"]))
    both = sum(1 for i in range(7) for j in range(7) if i != j and sum(F[y]["mv"][i, j] for y in calc.ALLW) > 0 and
               sum(F[y]["mv"][j, i] for y in calc.ALLW) > 0)
    C("C17 flows kept gross: moves run both ways between some pairs of schools", both >= 4, both)
    C("C18 units: a classroom is 24 pupils", "divided by 24" in pol)
    leg_grades = {r["GRADE"] for r in csv.DictReader(open(os.path.join(target, REC.F_LEGACY), encoding="utf-8",
                                                          newline=""), delimiter="|")}
    entry = {r["entry_type"] for r in rev}
    C("C19 codes: grades KG to 05 and 00 to 05; entry types 01 to 06", leg_grades == {"KG", "01", "02", "03", "04", "05"}
      and entry == {"01", "02", "03", "04", "05", "06"}, (leg_grades, entry))
    C("C20 integerisation: whole classrooms summing to 24", all(float(x).is_integer() for x in g["alloc"]) and
      int(g["alloc"].sum()) == 24)
    C("C21 clause scope: the regulation covers the Elementary Capacity program only",
      "governed by their own regulations" in pol)
    C("C22 forward window: births for 2026 and 2027 at the 2025 value or on the trend, the same schedule", all(
        placed(D, M["alloc"][k]) == GOLDEN_PLACED for k in ("births 2026-27 at 2025",
                                                            "births 2026-27 on the 2019-25 trend")))

    # ---- Ask A
    truthA = {}
    for z in K.ZONES:
        for b, gs in (("K-2", (0, 1, 2)), ("3-5", (3, 4, 5))):
            cls = [c for c in W.A.classes[z] if min(c["grades"]) in gs]
            truthA[(CODE[z], b)] = len({(c["room"], c.get("session", "FULL")) for c in cls})
    C("A1 Ask A golden equals the world's classes on October 1", truthA == A["golden"], (truthA, A["golden"]))
    C("A2 every class keeps its sections inside one band", A["one_band"])
    st = A["stops"]
    over = sum(1 for k2 in A["golden"] if st["sections running with pupils"][k2] > A["golden"][k2])
    C("A3 sections counted: every one of the 14 figures over", over == 14, over)
    off_t = sum(1 for k2 in A["golden"] if st["merged by teacher"][k2] != A["golden"][k2])
    C("A3' sections merged by teacher: every one of the 14 figures off", off_t == 14, off_t)
    off_r = sorted(k2 for k2 in A["golden"] if st["merged by room"][k2] != A["golden"][k2])
    C("A3'' sections merged by room: off at the three half-day kindergarten bands",
      off_r == sorted([(CODE[z], "K-2") for z in AS.HALF_DAY_K]), off_r)
    off_f = sum(1 for k2 in A["golden"] if st["every fall homeroom section"][k2] != A["golden"][k2])
    C("A3''' every fall homeroom section: every figure over", off_f == 14, off_f)
    C("A3'''' classes without the pupil test: off at North Fork 3-5 (the placeholder)",
      [k2 for k2 in A["golden"] if st["classes without the pupil test"][k2] != A["golden"][k2]] == [("6189", "3-5")])
    secs, lines = calc.read_sections(target)
    on1 = [l for l in lines if l["end_date"] == "2026-10-01" or l["start_date"] == "2026-10-01"]
    on1 += [s for s in secs if s["section_end"] == "2026-10-01" or s["section_start"] == "2026-10-01"]
    C("A4 nothing starts or ends on October 1 (boundary neutral)", not on1)
    ended = [s for s in secs if s["course_code"].startswith("HR-") and s["section_end"] == "2026-09-18"]
    C("A5 HA1 a homeroom section ended mid-September at Bluestem and Lincoln",
      sorted(s["school_code"] for s in ended) == sorted([CODE["B"], CODE["C"]]))
    C("A5' HA2 a placeholder homeroom section with no pupils at North Fork",
      [s["school_code"] for s in secs if s["room"] == "TBD"] == [CODE["F"]])
    combos = {}
    for s in secs:
        if s["course_code"].startswith("HR-"):
            combos.setdefault((s["school_code"], s["room"], s["session"]), []).append(s)
    shared_t = sum(1 for v in combos.values() if len(v) == 2 and v[0]["teacher_id"] == v[1]["teacher_id"] and
                   v[0]["course_code"] != v[1]["course_code"])
    cot = sum(1 for v in combos.values() if len(v) == 2 and v[0]["teacher_id"] != v[1]["teacher_id"] and
              v[0]["course_code"] == v[1]["course_code"] and v[0]["section_end"] == v[1]["section_end"])
    C("A6 A1 device: a combination class in each band at every school (14) and a co-taught class in every 3-5 band",
      shared_t == 14 and cot >= 11, (shared_t, cot))

    # ---- Ask B
    truthB = {K.SCHOOLS[z]["name"]: dict(children=len(W.P.truth[z]["kids"]), classrooms=len(W.P.truth[z]["rooms"]))
              for z in AS.PRESCHOOL}
    C("B1' Ask B golden equals the world's preschool rolls", truthB == B["golden"], (truthB, B["golden"]))
    sb = B["stops"]
    for k2 in ("platform child records and class names", "platform by state ID and room"):
        miss = sum(1 for s in B["sites"] for q in ("children", "classrooms") if sb[k2][s][q] != B["golden"][s][q])
        C("B2 stop %s: all 8 figures off" % k2, miss == 8, miss)
    for k2 in ("platform and roster added", "union with the itinerant children"):
        miss = sum(1 for s in B["sites"] if sb[k2][s]["children"] > B["golden"][s]["children"])
        C("B2' stop %s: children over at all 4 sites" % k2, miss == 4, miss)
    miss = sum(1 for s in B["sites"] if sb["union by state ID, class names as classrooms"][s]["classrooms"] !=
               B["golden"][s]["classrooms"])
    C("B2'' class names as classrooms with the roster's rooms: off at all 4 sites", miss == 4, miss)
    plat, roster = calc.read_preschool(target)
    dupe = {}
    for p in plat:
        dupe.setdefault(p["state_id"], set()).add(p["child_id"])
    C("B3 B2 device: a re-registered child under two child records at every site",
      len({p["site"] for p in plat if len(dupe[p["state_id"]]) > 1}) == 4)
    ecse_rooms = {x["Case school"] for x in roster if x["Service model"] == "ECSE classroom"}
    plat_rooms = {(p["site"], calc.ROOM.match(p["class_name"]).group(1)) for p in plat}
    C("B3' B1 device: a special education classroom at every site, none of its rooms on the platform",
      len(ecse_rooms) == 4 and not any((x["Case school"], x["Classroom"]) in plat_rooms for x in roster
                                       if x["Service model"] == "ECSE classroom"))
    C("B4 nothing on the platform starts or ends on October 1",
      not any(p["enroll_start"] == "2026-10-01" or p["enroll_end"] == "2026-10-01" for p in plat))

    # ---- the cracker sheet
    cracker_a = sum(1 for k2 in A["golden"] if st["sections running with pupils"][k2] == A["golden"][k2])
    cracker_b = sum(1 for s in B["sites"] for q in ("children", "classrooms")
                    if sb["platform child records and class names"][s][q] == B["golden"][s][q])
    C("CR the cracker's path misses every ask figure (Ask A 14, Ask B 8)", cracker_a == 0 and cracker_b == 0,
      (cracker_a, cracker_b))

    # ---- the separation line and the distractors
    tmp = tempfile.mkdtemp(prefix="t150sep")
    try:
        for f in MAIN_FILES:
            shutil.copy(os.path.join(target, f), tmp)
        M2 = calc.main_call(tmp, B=0)
        C("SEP1 the call from the main-call files alone is the golden (no ask file or distractor read)",
          placed(M2["D"], M2["golden"]) == GOLDEN_PLACED and np.allclose(M2["golden"]["exc"], g["exc"]))
        for f in ASK_FILES:
            shutil.copy(os.path.join(target, f), tmp)
        A2 = calc.ask_a(tmp, D.codes)
        B2 = calc.ask_b(tmp, D.names)
        C("SEP2 both asks unchanged with the distractors and the notes deleted", A2["golden"] == A["golden"] and
          B2["golden"] == B["golden"])
    finally:
        shutil.rmtree(tmp)
    main_rows = [r for r in rev if r["grade_level"] not in ("00", "01", "02", "03", "04", "05")]
    C("SEP3 zero device or hazard rows in the main call's population (no preschool or section rows in the membership)",
      not main_rows and not any(f in MAIN_FILES for f in ASK_FILES))
    meals = list(csv.DictReader(open(os.path.join(target, DOC.F_MEALS), encoding="utf-8")))
    C("SEP4 the meal eligibility counts tie to the October 1, 2026 count (a distractor that ranks nothing)", all(
        int(r["k5_pupils_oct1"]) == int(D.E[2026][D.codes.index(r["school_code"])].sum()) for r in meals))

    # ---- the realism sweep
    tot = [int(D.E[y].sum()) for y in D.years]
    C("RS1 no round district totals", all(t % 100 != 0 for t in tot), tot)
    C("RS2 rated capacities not round", all(c % 100 != 0 for c in D.cap) and sum(1 for c in D.cap if c % 10 == 0) <= 2,
      D.cap)
    caps_rows = W.cap_rows
    C("RS3 the capacity report's notes do not repeat", len({r[-1] for r in caps_rows if r[-1]}) ==
      len([r for r in caps_rows if r[-1]]))
    people = list(K.PEOPLE.values())
    alltext = " ".join(texts.values())
    C("RS4 the people named are the drawn personas", all(p in alltext for p in people[1:4] + [people[4]]) and
      not re.search(r"John Doe|Jane Doe|Acme|Company [A-Z]\b", alltext))
    bad_meta = []
    for f in files + ["golden/" + x for x in os.listdir(golden)]:
        p = os.path.join(target, f) if not f.startswith("golden/") else os.path.join(golden, f[7:])
        raw = open(p, "rb").read()
        if f.endswith(".xlsx"):
            z = zipfile.ZipFile(p)
            core = z.read("docProps/core.xml").decode()
            app = z.read("docProps/app.xml").decode()
            if "openpyxl" in core or "openpyxl" in app or "python" in (core + app).lower():
                bad_meta.append(f)
            m = re.search(r"<dcterms:modified[^>]*>([^<]+)<", core)
            when = dt.datetime.strptime(m.group(1), "%Y-%m-%dT%H:%M:%SZ").date()
            if not f.startswith("golden/") and when > ASOF:
                bad_meta.append((f, when))
        if f.endswith(".pdf") and (b"ReportLab" in raw or b"reportlab" in raw or b"D:2000" in raw):
            bad_meta.append(f)
        if f.endswith(".png") and (b"matplotlib" in raw.lower() or b"Software" in raw):
            bad_meta.append(f)
    C("RS5 containers scrubbed: no writer's name, no stamp after the pull (H1)", not bad_meta, bad_meta)
    dash = [f for f, t in texts.items() if DASH.search(t)]
    for p in [os.path.join(TASK, "submission.md"), os.path.join(TASK, "prompt.txt"), os.path.join(TASK, "DESIGN_NOTE.md")]:
        if os.path.exists(p) and DASH.search(open(p, encoding="utf-8").read()):
            dash.append(p)
    C("RS6 no en or em dash anywhere shipped or written", not dash, dash)
    same = sum(7 - len(set(D.E[y].sum(axis=1))) for y in D.years)
    C("RS7 two schools share a year's K-5 count at most once in 14 years", same <= 1, same)
    C("RS8 school names are not near a real school's distinctive word (generic only)", all(
        n.split()[0] in ("Fifth", "Lincoln", "Ortega", "Bluestem", "Prairie", "Kessler", "North") for n in names))

    # ---- goldens and the write-up
    figs = json.load(open(os.path.join(K.HERE, "figures.json"), encoding="utf-8"))
    wb = openpyxl.load_workbook(os.path.join(golden, DL.F_BOOK), data_only=True)
    ws = wb["Placement"]
    book = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[1] in D.codes:
            book[row[1]] = (int(row[6]), int(row[8]))
    C("GD1 the workbook's needs and placements equal the golden",
      book == {c: (GOLDEN_NEED[c], GOLDEN_PLACED[c]) for c in D.codes}, book)
    hw = wb["Homeroom classes"]
    hb = {row[0]: (row[1], row[2]) for row in hw.iter_rows(min_row=4, values_only=True) if row[0] in D.names}
    C("GD2 the workbook's homeroom table equals Ask A",
      hb == {D.names[i]: (A["golden"][(c, "K-2")], A["golden"][(c, "3-5")]) for i, c in enumerate(D.codes)}, hb)
    pw = wb["Preschool"]
    pb = {row[0]: (row[1], row[2]) for row in pw.iter_rows(min_row=4, values_only=True) if row[0] in B["sites"]}
    C("GD3 the workbook's preschool table equals Ask B",
      pb == {s: (B["golden"][s]["children"], B["golden"][s]["classrooms"]) for s in B["sites"]}, pb)
    C("GD4 figures.json carries the call", {DL.SHORT[c]: int(v) for c, v in GOLDEN_PLACED.items()} ==
      {k2: int(v) for k2, v in figs["call"].items()})
    C("GD5 two goldens, both named in the prompt", sorted(os.listdir(golden)) == sorted([DL.F_BOOK, DL.F_CHART]))
    sub = open(os.path.join(TASK, "submission.md"), encoding="utf-8").read()
    blocks = re.findall(r"^## (.+)$", sub, re.M)
    C("GD6 the write-up carries exactly the five blocks", blocks == ["Tags", "1. Final Recommendation",
                                                                     "2. Critical Components",
                                                                     "3. Step-by-Step Solution",
                                                                     "4. Deliverable Answers"], blocks)
    steps = re.findall(r"^\d+\. ", sub.split("## 3.")[1].split("## 4.")[0], re.M)
    C("GD7 at most eight steps", len(steps) <= 8, len(steps))
    named = set(re.findall(r"[\w.-]+\.(?:csv|txt|xlsx|pdf)", sub.split("## 4.")[0]))
    C("GD8 every file the steps name ships", named <= set(files), named - set(files))

    # ---- the prompt
    prompt = open(os.path.join(TASK, "prompt.txt"), encoding="utf-8").read().strip()
    C("P1 the prompt is one paragraph under 1,000 characters naming both deliverables",
      len(prompt) < 1000 and "\n" not in prompt and DL.F_BOOK in prompt and DL.F_CHART in prompt, len(prompt))
    C("P2 the prompt names no method, ratio, rate or flow", not method.search(prompt))
    C.report()
    return C


def ng_share(target, y):
    """The not-geocoded row's share of the areas' births in year y (to spread it pro rata)."""
    rows = list(csv.DictReader(open(os.path.join(target, REC.F_BIRTHS), encoding="utf-8")))
    areas = sum(int(r["births"]) for r in rows if int(r["birth_year"]) == y and r["attendance_area"] not in
                ("OUTSIDE RE-6", "NOT GEOCODED"))
    ng = sum(int(r["births"]) for r in rows if int(r["birth_year"]) == y and r["attendance_area"] == "NOT GEOCODED")
    return ng / areas if areas else 0.0
