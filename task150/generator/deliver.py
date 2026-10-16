"""The golden deliverables, computed from the shipped pack only (calc): classroom_placement_2028-33.xlsx (the placement
schedule, the projection behind it, and the two October 1, 2026 tables) and enrollment_vs_capacity_2026-33.png (one
panel per school), with generator/figures.json and submission.md written from the same objects."""
from __future__ import annotations

import datetime as dt
import json
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt          # noqa: E402
import numpy as np                       # noqa: E402
import openpyxl                          # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side   # noqa: E402

import calc                              # noqa: E402
import common as K                       # noqa: E402

F_BOOK = "classroom_placement_2028-33.xlsx"
F_CHART = "enrollment_vs_capacity_2026-33.png"
STAMP = dt.datetime(2026, 10, 15, 10, 20, 0)     # the planning office's save, six days after the pull
AUTHOR = K.PEOPLE["planning"]
SHORT = {"6111": "Fifth Street", "6116": "Lincoln", "6142": "Ortega", "6158": "Bluestem", "6163": "Prairie Rose",
         "6175": "Kessler", "6189": "North Fork"}


def compute(target):
    M = calc.main_call(target)
    D = M["D"]
    return dict(M=M, D=D, A=calc.ask_a(target, D.codes), B=calc.ask_b(target, D.names))


def n0(v):
    return "{:,}".format(int(round(float(v))))


def pct1(v):
    return "%.1f" % float(v)


def facts(R):
    """The in-sample flow facts and the back-test, as the write-up states them."""
    D, F = R["D"], R["M"]["F"]
    ix = {c: i for i, c in enumerate(D.codes)}
    P, E, A, B, Kd = ix["6111"], ix["6163"], ix["6142"], ix["6158"], ix["6175"]
    arr = {c: [F[y]["arr"][ix[c]].sum() for y in calc.ALLW] for c in D.codes}
    base = {c: [F[y]["base"][ix[c]] for y in calc.ALLW] for c in D.codes}
    a2b = [F[y]["mv"][A, B] for y in calc.ALLW]
    a2b_rate = [100.0 * F[y]["mv"][A, B] / F[y]["base"][A] for y in calc.ALLW]
    a2d = [F[y]["mv"][A, Kd] for y in calc.ALLW]
    a2d_rate = [100.0 * F[y]["mv"][A, Kd] / F[y]["base"][A] for y in calc.ALLW]
    bt = R["M"]["backtest"]
    g = bt["gross-flow"]["last"]
    worst = int(np.argmax(np.abs(g)))
    return dict(arr=arr, base=base, a2b=a2b, a2b_rate=a2b_rate, a2d=a2d, a2d_rate=a2d_rate, bt=bt, worst=worst,
                P=P, E=E, A=A, B=B)


def figures(R):
    D, M = R["D"], R["M"]
    g, d = M["golden"], M["default"]
    fc = facts(R)
    bt = fc["bt"]
    rows = sum(D.rows.values())
    f = {
        "call": {SHORT[c]: str(int(g["alloc"][i])) for i, c in enumerate(D.codes)},
        "classrooms_total": "24", "site_cap": "8", "seats_per_classroom": "24",
        "critical_components": {
            "classroom_need": {SHORT[c]: str(int(g["need"][i])) for i, c in enumerate(D.codes)},
            "classroom_need_total": str(int(g["need"].sum())),
            "order_above_capacity": [SHORT[D.codes[i]] for i in g["order"]],
        },
        "steps": {
            "pupil_rows": n0(rows), "snapshots": str(len(D.years)),
            "first_year": str(D.years[0]), "last_year": str(D.years[-1]), "transitions": str(len(calc.ALLW)),
            "fifth_street_arrivals_low": n0(min(fc["arr"]["6111"])), "fifth_street_arrivals_high": n0(max(fc["arr"]["6111"])),
            "fifth_street_base_first": n0(fc["base"]["6111"][0]), "fifth_street_base_last": n0(fc["base"]["6111"][-1]),
            "prairie_rose_arrivals_low": n0(min(fc["arr"]["6163"])), "prairie_rose_arrivals_high": n0(max(fc["arr"]["6163"])),
            "prairie_rose_base_first": n0(fc["base"]["6163"][0]), "prairie_rose_base_last": n0(fc["base"]["6163"][-1]),
            "ortega_to_bluestem_low": n0(min(fc["a2b"])), "ortega_to_bluestem_high": n0(max(fc["a2b"])),
            "ortega_to_bluestem_rate_low": pct1(min(fc["a2b_rate"])),
            "ortega_to_bluestem_rate_high": pct1(max(fc["a2b_rate"])),
            "ortega_rate_per": "100",
            "backtest_origin": "2019",
            "backtest_gross_worst_abs": pct1(abs(bt["gross-flow"]["last"][fc["worst"]])),
            "backtest_gross_worst_school": SHORT[D.codes[fc["worst"]]],
            "backtest_gpr5_fifth_street": pct1(bt["five-year progression ratios"]["last"][fc["P"]]),
            "backtest_gpr5_prairie_rose": pct1(bt["five-year progression ratios"]["last"][fc["E"]]),
            "backtest_gpr5_fifth_street_abs": pct1(abs(bt["five-year progression ratios"]["last"][fc["P"]])),
            "backtest_gpr5_prairie_rose_abs": pct1(abs(bt["five-year progression ratios"]["last"][fc["E"]])),
            "backtest_gross_fifth_street": pct1(bt["gross-flow"]["last"][fc["P"]]),
            "backtest_gross_prairie_rose": pct1(bt["gross-flow"]["last"][fc["E"]]),
            "excess": {SHORT[c]: n0(g["exc"][i]) for i, c in enumerate(D.codes)},
            "ortega_below": n0(-g["exc"][D.codes.index("6142")]),
            "five_year_average": {SHORT[c]: n0(g["avg"][i]) for i, c in enumerate(D.codes)},
            "rated_capacity": {SHORT[c]: n0(D.cap[i]) for i, c in enumerate(D.codes)},
            "october_2026": {SHORT[c]: n0(D.E[2026][i].sum()) for i, c in enumerate(D.codes)},
            "kessler_remaining": str(int(g["alloc"][D.codes.index("6175")])),
            "births_2026_27_basis": "2023-25",
        },
        "default_path": {
            "placement": {SHORT[c]: str(int(d["alloc"][i])) for i, c in enumerate(D.codes)},
            "need": {SHORT[c]: str(int(d["need"][i])) for i, c in enumerate(D.codes)},
            "window": "2021-22 to 2025-26",
        },
        "rivals": {k: {SHORT[c]: str(int(M["alloc"][k]["alloc"][i])) for i, c in enumerate(D.codes)}
                   for k in M["rivals"]},
        "projection": {SHORT[c]: {str(y): n0(M["P"][calc.GOLDEN][y][i].sum()) for y in calc.FUT}
                       for i, c in enumerate(D.codes)},
        "ask_a_homeroom_classes": [{"school": D.names[i], "k2": str(R["A"]["golden"][(c, "K-2")]),
                                    "g35": str(R["A"]["golden"][(c, "3-5")])} for i, c in enumerate(D.codes)],
        "ask_a_district": {"k2": str(sum(R["A"]["golden"][(c, "K-2")] for c in D.codes)),
                           "g35": str(sum(R["A"]["golden"][(c, "3-5")] for c in D.codes))},
        "ask_b_preschool": [{"site": s, "children": str(R["B"]["golden"][s]["children"]),
                             "classrooms": str(R["B"]["golden"][s]["classrooms"])} for s in R["B"]["sites"]],
        "ask_b_total": {"children": str(sum(v["children"] for v in R["B"]["golden"].values())),
                        "classrooms": str(sum(v["classrooms"] for v in R["B"]["golden"].values()))},
        "chart_years": {"history_from": "2013", "history_to": "2026", "projection_to": "2032", "placed_from": "2028"},
    }
    return f


# ---------------------------------------------------------------------------------------------------- workbook
THIN = Side(style="thin", color="999999")
HEAD_FILL = PatternFill("solid", fgColor="DCE6F0")
PLACED_FILL = PatternFill("solid", fgColor="FCE9D9")


def _header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="bottom", horizontal="left" if c == 1 else "right")
        cell.border = Border(bottom=THIN)


def write_book(path, R):
    D, M = R["D"], R["M"]
    g = M["golden"]
    P = M["P"][calc.GOLDEN]
    fc = facts(R)
    bt = fc["bt"]
    placed = [i for i in g["order"] if g["alloc"][i] > 0]
    title = ("Elementary Capacity program: the 24 classrooms placed at %s" %
             ", ".join("%s (%d)" % (SHORT[D.codes[i]], g["alloc"][i]) for i in placed))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Placement"
    ws.append([title])
    ws.append(["Placement schedule under Board Policy FEA-R2 for the school years 2028-29 to 2032-33, from the October 1, "
               "2026 count. Planning and Enrollment Office, %s." % STAMP.strftime("%m/%d/%Y")])
    ws.append([])
    hdr = ["School", "State code", "K-5 enrollment, Oct 1, 2026", "Rated capacity",
           "Five-year average projected enrollment, 2028-29 to 2032-33", "Above rated capacity (pupils)",
           "Classroom need", "Order", "Classrooms placed"]
    ws.append(hdr)
    _header(ws, 4, len(hdr))
    order_no = {i: k + 1 for k, i in enumerate(g["order"])}
    for i, c in enumerate(D.codes):
        ws.append([D.names[i], c, int(D.E[2026][i].sum()), int(D.cap[i]), round(float(g["avg"][i])),
                   round(float(g["exc"][i])), int(g["need"][i]), order_no.get(i, ""), int(g["alloc"][i])])
        if g["alloc"][i] > 0:
            ws.cell(row=ws.max_row, column=9).fill = PLACED_FILL
    tot = ws.max_row + 1
    ws.append(["District", "", int(sum(D.E[2026][i].sum() for i in range(D.NS))), int(D.cap.sum()),
               round(float(g["avg"].sum())), "", int(g["need"].sum()), "", int(g["alloc"].sum())])
    for c in range(1, len(hdr) + 1):
        ws.cell(row=tot, column=c).font = Font(bold=True)
        ws.cell(row=tot, column=c).border = Border(top=THIN)
    ws.append([])
    kz = D.codes.index("6175")
    notes = [
        "Classroom need (FEA-R2 2(e)): the unrounded five-year average less rated capacity, divided by 24 and rounded up; "
        "none where the average does not exceed rated capacity.",
        "Placement (FEA-R2 3): in descending order of the amount above rated capacity, each school its need or eight, "
        "whichever is fewer, until the 24 are placed. %s's need is %d; the %d classrooms left when its turn came are "
        "placed there, and North Fork's need of %d is not reached." % (
            SHORT["6175"], g["need"][kz], g["alloc"][kz], g["need"][D.codes.index("6189")]),
        "Rated capacity: facilities_capacity_report_2026-09.xlsx (permanent classrooms; portables not counted).",
    ]
    for t in notes:
        ws.append([t])
    widths = [26, 9, 13, 11, 20, 13, 11, 8, 12]
    for k, w in enumerate(widths):
        ws.column_dimensions[chr(65 + k)].width = w
    ws.row_dimensions[4].height = 58
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"].font = Font(italic=True, size=9)
    for row in ws.iter_rows(min_row=5, max_row=tot):
        for cell in row[2:]:
            cell.number_format = "#,##0"
            cell.alignment = Alignment(horizontal="right")
    for r in range(tot + 2, ws.max_row + 1):
        ws.cell(row=r, column=1).font = Font(size=9)
    ws.freeze_panes = "B5"
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1

    # the projection
    pj = wb.create_sheet("Projection")
    pj.append(["Projected October K-5 enrollment by school, 2027 to 2032 (October 1, 2026 count shown for reference)"])
    pj.append([])
    yrs = [2026] + calc.FUT
    hdr2 = ["School"] + ["Oct %d%s" % (y, " (count)" if y == 2026 else "") for y in yrs] + ["Five-year average, 2028-32"]
    pj.append(hdr2)
    _header(pj, 3, len(hdr2))
    for i, c in enumerate(D.codes):
        pj.append([D.names[i]] + [round(float(P[y][i].sum())) for y in yrs] + [round(float(g["avg"][i]))])
    tr = pj.max_row + 1
    pj.append(["District"] + [round(float(P[y].sum())) for y in yrs] + [round(float(g["avg"].sum()))])
    for c in range(1, len(hdr2) + 1):
        pj.cell(row=tr, column=c).font = Font(bold=True)
        pj.cell(row=tr, column=c).border = Border(top=THIN)
    for row in pj.iter_rows(min_row=4, max_row=tr):
        for cell in row[1:]:
            cell.number_format = "#,##0"
    pj.append([])
    pj.append(["Projected from the October 1, 2026 count, grade by grade; see Notes for the method and its back-test. "
               "Kindergartens for 2031 and 2032 rest on births for 2026 and 2027, taken at the 2023-25 average."])
    pj.cell(row=pj.max_row, column=1).font = Font(size=9)
    pj["A1"].font = Font(bold=True, size=12)
    pj.column_dimensions["A"].width = 26
    for k in range(2, len(hdr2) + 1):
        pj.column_dimensions[chr(64 + k)].width = 11
    pj.column_dimensions[chr(64 + len(hdr2))].width = 14
    pj.row_dimensions[3].height = 32
    pj.freeze_panes = "B4"

    # Ask A
    ha = wb.create_sheet("Homeroom classes")
    ha.append(["Homeroom classes meeting on October 1, 2026"])
    ha.append([])
    hdr3 = ["School", "Kindergarten to grade 2", "Grades 3 to 5"]
    ha.append(hdr3)
    _header(ha, 3, 3)
    for i, c in enumerate(D.codes):
        ha.append([D.names[i], R["A"]["golden"][(c, "K-2")], R["A"]["golden"][(c, "3-5")]])
    tr = ha.max_row + 1
    ha.append(["District", sum(R["A"]["golden"][(c, "K-2")] for c in D.codes),
               sum(R["A"]["golden"][(c, "3-5")] for c in D.codes)])
    for c in range(1, 4):
        ha.cell(row=tr, column=c).font = Font(bold=True)
        ha.cell(row=tr, column=c).border = Border(top=THIN)
    ha.append([])
    ha.append(["One class for each room and session in use on October 1: a combination class's grade sections and a "
               "co-taught class's two sections count once, a half-day kindergarten's morning and afternoon classes twice. "
               "Homeroom sections (HR-K to HR-5) running on October 1 with at least one pupil enrolled that day "
               "(sections_fall2026.csv, section_enrollment_fall2026.csv)."])
    ha.cell(row=ha.max_row, column=1).font = Font(size=9)
    ha.cell(row=ha.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ha.merge_cells(start_row=ha.max_row, start_column=1, end_row=ha.max_row, end_column=3)
    ha.row_dimensions[ha.max_row].height = 64
    ha["A1"].font = Font(bold=True, size=12)
    for k, w in enumerate([26, 14, 14]):
        ha.column_dimensions[chr(65 + k)].width = w
    ha.freeze_panes = "B4"

    # Ask B
    pk = wb.create_sheet("Preschool")
    pk.append(["Preschool at the elementary schools on October 1, 2026"])
    pk.append([])
    hdr4 = ["School", "Children enrolled", "Classrooms used"]
    pk.append(hdr4)
    _header(pk, 3, 3)
    for s in R["B"]["sites"]:
        pk.append([s, R["B"]["golden"][s]["children"], R["B"]["golden"][s]["classrooms"]])
    tr = pk.max_row + 1
    pk.append(["Total", sum(v["children"] for v in R["B"]["golden"].values()),
               sum(v["classrooms"] for v in R["B"]["golden"].values())])
    for c in range(1, 4):
        pk.cell(row=tr, column=c).font = Font(bold=True)
        pk.cell(row=tr, column=c).border = Border(top=THIN)
    pk.append([])
    pk.append(["Children: each child once (state ID), from the registration platform's enrollments open on October 1 "
               "(preschool_enrollments_2026-27.csv) and the special education roster's children placed in the school's "
               "special education and inclusive preschool classrooms (ecse_roster_2026-10-01.xlsx); children served "
               "itinerantly are not in a class. Classrooms: the rooms those classes meet in, morning and afternoon "
               "classes in one room counted once."])
    pk.cell(row=pk.max_row, column=1).font = Font(size=9)
    pk.cell(row=pk.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    pk.merge_cells(start_row=pk.max_row, start_column=1, end_row=pk.max_row, end_column=3)
    pk.row_dimensions[pk.max_row].height = 78
    pk["A1"].font = Font(bold=True, size=12)
    for k, w in enumerate([26, 14, 14]):
        pk.column_dimensions[chr(65 + k)].width = w
    pk.freeze_panes = "B4"

    # Notes
    nt = wb.create_sheet("Notes")
    gl = bt["gross-flow"]["last"]
    gp = bt["five-year progression ratios"]["last"]
    lines = [
        ("Notes on the projection", True),
        ("", False),
        ("Projection. Each school's October K-5 enrollment is projected grade by grade from the October 1, 2026 count. "
         "Between one October and the next a pupil repeats the grade, moves up a grade at the school, moves to another "
         "district school or leaves the district. Repeats are rated on the pupils in that grade at the school; moves "
         "(by destination school) and exits on all of the school's K-4 pupils, the school the pupil leaves. Pupils new "
         "to the district in grades 1 to 5 are added at each school's average count per grade. Kindergartners new to "
         "the district are each area's resident births five years earlier times the school's ratio of new "
         "kindergartners to those births. Every rate, count and ratio is taken over the %s transitions from %d-%02d "
         "to %d-%02d in the linked October records; births for 2026 and 2027 are taken at the 2023-25 average." % (
             len(calc.ALLW), calc.ALLW[0] - 1, calc.ALLW[0] % 100, calc.ALLW[-1] - 1, calc.ALLW[-1] % 100), False),
        ("", False),
        ("Why flow by flow. Fifth Street took %s to %s pupils a year from outside the district into grades 1 to 5, and "
         "Prairie Rose %s to %s, in every year from 2014 while Fifth Street's K-4 pupils fell from %s to %s and Prairie "
         "Rose's rose from %s to %s; Ortega sent Bluestem %s to %s pupils a year, %s to %s for every 100 Ortega K-4 "
         "pupils. A grade progression ratio strikes those arrivals and transfers on the receiving class, so it carries "
         "the past sizes forward. Fitted on the transitions to 2019 and run from the 2019 count with the births on "
         "file, this projection misses each school's 2026 count by %s per cent or less (%s the largest); five-year "
         "progression ratios come out %s per cent low at Fifth Street and %s per cent high at Prairie Rose." % (
             n0(min(fc["arr"]["6111"])), n0(max(fc["arr"]["6111"])), n0(min(fc["arr"]["6163"])),
             n0(max(fc["arr"]["6163"])), n0(fc["base"]["6111"][0]), n0(fc["base"]["6111"][-1]),
             n0(fc["base"]["6163"][0]), n0(fc["base"]["6163"][-1]), n0(min(fc["a2b"])), n0(max(fc["a2b"])),
             pct1(min(fc["a2b_rate"])), pct1(max(fc["a2b_rate"])), pct1(abs(gl[fc["worst"]])),
             SHORT[D.codes[fc["worst"]]], pct1(abs(gp[fc["P"]])), pct1(abs(gp[fc["E"]]))), False),
        ("", False),
        ("Records. October membership snapshots 2013-14 to 2018-19 (legacy system) and 2019-20 to 2026-27 (current "
         "system), linked through the 2019 student number crosswalk; Southeast Elementary is Kessler Elementary (codes "
         "105 and 6175). Births: Brandt County resident births by attendance area, 2008 to 2025 (2025 provisional); "
         "births outside the district and not geocoded are not in any area.", False),
        ("", False),
        ("Rounding. Figures are whole numbers. Classroom need is computed from the unrounded five-year average.", False),
        ("Prepared by %s, Planning and Enrollment Office, %s, for the Board's November 12, 2026 meeting." % (
            AUTHOR, STAMP.strftime("%m/%d/%Y")), False),
    ]
    for t, bold in lines:
        nt.append([t])
        cell = nt.cell(row=nt.max_row, column=1)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = Font(bold=bold, size=12 if bold else 10)
        if t and not bold:
            nt.row_dimensions[nt.max_row].height = max(15, 15 * (len(t) // 105 + 1))
    nt.column_dimensions["A"].width = 120
    K.save_workbook(wb, path, STAMP, AUTHOR, "Elementary classroom placement 2028-33")


# ---------------------------------------------------------------------------------------------------- chart
def write_chart(path, R):
    D, M = R["D"], R["M"]
    g = M["golden"]
    P = M["P"][calc.GOLDEN]
    hist = list(range(D.years[0], 2027))
    fut = [2026] + calc.FUT
    fig, axes = plt.subplots(2, 4, figsize=(15, 7.6), dpi=110)
    order = [D.codes.index(c) for c in ("6163", "6116", "6111", "6158", "6175", "6189", "6142")]
    for k, i in enumerate(order):
        ax = axes.flat[k]
        h = [D.E[y][i].sum() for y in hist]
        pr = [D.E[2026][i].sum()] + [P[y][i].sum() for y in calc.FUT]
        ax.axvspan(2027.5, 2032.5, color="#F2F2F2", zorder=0)
        ax.plot(hist, h, color="#1F3B5C", lw=2.0, label="October count")
        col = "#C0502E" if g["alloc"][i] > 0 else "#5B7A99"
        ax.plot(fut, pr, color=col, lw=2.0, ls="--", label="Projection")
        cap = D.cap[i]
        ax.axhline(cap, color="#7F7F7F", lw=1.2, label="Rated capacity")
        if g["alloc"][i] > 0:
            ax.step([2027.5, 2032.5], [cap + 24 * g["alloc"][i]] * 2, where="post", color="#C0502E", lw=1.2, ls=":",
                    label="Capacity with placed classrooms")
        pts = list(zip(hist, h)) + list(zip(fut, pr))
        clear_right = not any(x >= 2027.5 and cap - 26 <= v <= cap for x, v in pts)
        clear_left = not any(x <= 2017.5 and cap - 26 <= v <= cap for x, v in pts)
        if clear_right:
            ax.text(2032.4, cap - 5, "rated %s" % n0(cap), fontsize=8, color="#555555", va="top", ha="right")
        elif clear_left:
            ax.text(2012.8, cap - 5, "rated %s" % n0(cap), fontsize=8, color="#555555", va="top", ha="left")
        else:
            ax.text(2032.4, cap + 5, "rated %s" % n0(cap), fontsize=8, color="#555555", va="bottom", ha="right")
        need = int(g["need"][i])
        ax.set_title("%s: %s placed (need %d)" % (SHORT[D.codes[i]], "%d" % g["alloc"][i] if g["alloc"][i] else "none",
                                                   need), fontsize=10.5, loc="left")
        lo = min(min(h), min(pr), cap) - 40
        hi = max(max(h), max(pr), cap + 24 * g["alloc"][i]) + 40
        ax.set_ylim(lo, hi)
        ax.set_xlim(2012.5, 2032.5)
        ax.set_xticks([2013, 2017, 2021, 2025, 2029, 2032])
        ax.tick_params(labelsize=8.5)
        ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, p: "{:,.0f}".format(v)))
        ax.grid(axis="y", color="#E5E5E5", lw=0.6)
        for s in ("top", "right"):
            ax.spines[s].set_visible(False)
        if k % 4 == 0:
            ax.set_ylabel("K-5 pupils, October count", fontsize=9)
    leg = axes.flat[7]
    leg.axis("off")
    hs, ls = axes.flat[0].get_legend_handles_labels()
    leg.legend(hs, ls, loc="upper left", fontsize=9.5, frameon=False)
    leg.text(0.0, 0.42, "Shaded: the school years 2028-29 to 2032-33\nthe new classrooms serve.\n\nOctober counts 2013 to "
             "2026 from the\nmembership snapshots; projection from the\nOctober 1, 2026 count. Rated capacity from\nthe "
             "September 2026 capacity report.", fontsize=8.5, va="top", color="#333333", transform=leg.transAxes)
    fig.suptitle("The 24 bond classrooms go to Prairie Rose, Lincoln, Fifth Street, Bluestem and Kessler: projected "
                 "October K-5 enrollment against rated capacity", fontsize=12.5, x=0.01, ha="left", fontweight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.95))
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)


# ---------------------------------------------------------------------------------------------------- write-up
def submission(R, f):
    D = R["D"]
    s = f["steps"]
    c = f["call"]
    nd = f["critical_components"]["classroom_need"]
    dp = f["default_path"]["placement"]
    placed = [SHORT[D.codes[i]] for i in R["M"]["golden"]["order"] if R["M"]["golden"]["alloc"][i] > 0]
    out = []
    out.append("## Tags\n")
    out.append("**Domain:** Sociology & Demography (population estimation and demographic projection: each attendance "
               "area's October K-5 pupils projected grade by grade over the school years 2028-29 to 2032-33, with exits "
               "and moves between areas rated on the area a pupil leaves and arrivals from outside the district as "
               "counts).")
    out.append("**Analytical objective:** Forecasting & Predictive Modeling.\n")
    out.append("## 1. Final Recommendation\n")
    out.append("**Place the 24 classrooms at %s, and none at %s.**" % (
        ", ".join("%s %s" % (p, c[p]) for p in placed[:-1]) + " and %s %s" % (placed[-1], c[placed[-1]]),
        " or ".join(x for x in ("Ortega", "North Fork") if c[x] == "0")))
    out.append("Projected flow by flow, the five schools need %s, %s, %s, %s and %s classrooms (%s with North Fork's %s), so "
               "the 24 run out at Kessler, which is placed the %s that remain. Not Bluestem %s, Fifth Street %s, Lincoln "
               "%s and Prairie Rose %s with none at Kessler, which five-year grade progression ratios give by striking "
               "Fifth Street's and Prairie Rose's steady stream of arrivals on classes of the past size and Ortega's "
               "transfers on Bluestem's and Kessler's classes; run from 2019, those ratios came out %s per cent low at "
               "Fifth Street and %s per cent high at Prairie Rose by 2026." % (
                   nd["Prairie Rose"], nd["Lincoln"], nd["Fifth Street"], nd["Bluestem"], nd["Kessler"],
                   f["critical_components"]["classroom_need_total"], nd["North Fork"], s["kessler_remaining"],
                   dp["Bluestem"], dp["Fifth Street"], dp["Lincoln"], dp["Prairie Rose"],
                   s["backtest_gpr5_fifth_street_abs"], s["backtest_gpr5_prairie_rose_abs"]))
    out.append("")
    out.append("## 2. Critical Components\n")
    out.append("1. Classroom need: Prairie Rose **%s**, Lincoln **%s**, Fifth Street **%s**, Bluestem **%s**, Kessler "
               "**%s**, North Fork **%s**, Ortega **%s** (**%s** in all against the program's 24)" % (
                   nd["Prairie Rose"], nd["Lincoln"], nd["Fifth Street"], nd["Bluestem"], nd["Kessler"],
                   nd["North Fork"], nd["Ortega"], f["critical_components"]["classroom_need_total"]))
    out.append("2. Order above rated capacity: %s; Ortega's five-year average is below its rated capacity" %
               ", ".join(f["critical_components"]["order_above_capacity"]))
    out.append("")
    out.append("## 3. Step-by-Step Solution\n")
    out.append("1. Linked every pupil's October records from membership_oct2013-oct2018.txt and "
               "october_membership_2019-2026.csv through student_number_crosswalk_2019.csv and school_codes.csv "
               "(Southeast is Kessler): %s pupil rows on %s count days." % (s["pupil_rows"], s["snapshots"]))
    out.append("2. Counted each of the %s transitions from one October to the next by the school a pupil left (exits, "
               "moves by destination, repeats) and the pupils new to the roll: Fifth Street took %s to %s arrivals into "
               "grades 1 to 5 a year and Prairie Rose %s to %s while their K-4 pupils went from %s to %s and from %s to "
               "%s, and Ortega sent Bluestem %s to %s a year, %s to %s per 100 Ortega K-4 pupils." % (
                   s["transitions"], s["fifth_street_arrivals_low"], s["fifth_street_arrivals_high"],
                   s["prairie_rose_arrivals_low"], s["prairie_rose_arrivals_high"], s["fifth_street_base_first"],
                   s["fifth_street_base_last"], s["prairie_rose_base_first"], s["prairie_rose_base_last"],
                   s["ortega_to_bluestem_low"], s["ortega_to_bluestem_high"], s["ortega_to_bluestem_rate_low"],
                   s["ortega_to_bluestem_rate_high"]))
    out.append("3. Projected each grade with repeats rated on that grade's pupils, moves and exits rated on the K-4 "
               "pupils of the school left, arrivals at each school's average count per grade, and kindergartners as "
               "new kindergartners per resident birth five years earlier (brandt_county_births_by_attendance_area_"
               "2008-2025.csv, births for 2026 and 2027 at the %s average), all over the %s transitions." % (
                   s["births_2026_27_basis"], s["transitions"]))
    out.append("4. Back-tested from %s (fitted on the transitions to %s, run from the %s count with the births on file): "
               "the projection misses each school's 2026 count by %s per cent or less, while five-year progression "
               "ratios come out %s per cent low at Fifth Street and %s per cent high at Prairie Rose." % (
                   s["backtest_origin"], s["backtest_origin"], s["backtest_origin"], s["backtest_gross_worst_abs"],
                   s["backtest_gpr5_fifth_street_abs"], s["backtest_gpr5_prairie_rose_abs"]))
    out.append("5. Rolled every school from the October 1, 2026 count to 2032 and averaged 2028 to 2032 against rated "
               "capacity in facilities_capacity_report_2026-09.xlsx: above capacity by %s pupils at Prairie Rose, %s at "
               "Lincoln, %s at Fifth Street, %s at Bluestem, %s at Kessler and %s at North Fork, and %s below at "
               "Ortega." % (
                   s["excess"]["Prairie Rose"], s["excess"]["Lincoln"], s["excess"]["Fifth Street"],
                   s["excess"]["Bluestem"], s["excess"]["Kessler"], s["excess"]["North Fork"],
                   s["ortega_below"]))
    out.append("6. Classroom need under BP_FEA-R2_elementary_classroom_placement.pdf 2(e) (excess over 24, rounded up): "
               "%s in all." % f["critical_components"]["classroom_need_total"])
    out.append("7. Placed under FEA-R2 3 in that order, the need or eight until 24: Prairie Rose %s, Lincoln %s, Fifth "
               "Street %s, Bluestem %s, Kessler %s." % (c["Prairie Rose"], c["Lincoln"], c["Fifth Street"],
                                                         c["Bluestem"], c["Kessler"]))
    out.append("")
    out.append("## 4. Deliverable Answers\n")
    out.append("### classroom_placement_2028-33.xlsx\n")
    out.append("1. Classrooms each school's projected enrollment calls for (its classroom need) and classrooms placed:")
    for i, code in enumerate(D.codes):
        nm = SHORT[code]
        out.append("   - %s: need %s, placed %s" % (D.names[i], nd[nm], c[nm]))
    out.append("2. Homeroom classes meeting on October 1, 2026:")
    for r in f["ask_a_homeroom_classes"]:
        out.append("   - %s: %s in kindergarten to grade 2, %s in grades 3 to 5" % (r["school"], r["k2"], r["g35"]))
    out.append("3. Preschool on October 1, 2026:")
    for r in f["ask_b_preschool"]:
        out.append("   - %s: %s children enrolled, %s classrooms" % (r["site"], r["children"], r["classrooms"]))
    out.append("")
    out.append("### enrollment_vs_capacity_2026-33.png\n")
    out.append("1. One panel per school: October K-5 enrollment %s to %s from the membership snapshots, the projection "
               "to %s, rated capacity across the panel and capacity with the placed classrooms from %s." % (
                   f["chart_years"]["history_from"], f["chart_years"]["history_to"], f["chart_years"]["projection_to"],
                   f["chart_years"]["placed_from"]))
    return "\n".join(out) + "\n"


def main(target, golden, figures_path, submission_path, R=None):
    R = R or compute(target)
    os.makedirs(golden, exist_ok=True)
    write_book(os.path.join(golden, F_BOOK), R)
    write_chart(os.path.join(golden, F_CHART), R)
    f = figures(R)
    with open(figures_path, "w", encoding="utf-8") as fh:
        json.dump(f, fh, indent=1, ensure_ascii=False)
        fh.write("\n")
    with open(submission_path, "w", encoding="utf-8") as fh:
        fh.write(submission(R, f))
    return R, f
