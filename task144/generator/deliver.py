"""The goldens, from the shipped pack only: sluis_spotlight_2026-10.xlsx (the pick with each shortlisted repository's
expected 90-day external stars if spotlighted, the April to September 2026 issue report and the nominations) and
shortlist_forecast_2026-10.png; then generator/figures.json, every figure as submission.md states it."""
from __future__ import annotations

import datetime as dt
import json
import math
import os

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import calc as C
import common as K
import postprocess as PP

INK = "1D2733"
GREY = "5A6470"
HEAD = "DCE7EB"
PICK = "FFF1CC"
THIN = Side(style="thin", color="B4BEC6")


def k(x):
    """Stars to the nearest thousand."""
    return K.round_to(x, 1000)


def fmt(n):
    return "{:,}".format(int(n))


def next_names(G):
    """Every candidate after the pick whose expected stars round to the runner-up's thousand: the order inside one
    thousand is not a finding at this rounding."""
    return [r["name"] for r in G["rows"][1:] if k(r["total"]) == k(G["runner"]["total"])]


def compute(target):
    B = C.load(target)
    R = C.rungs(B)
    A = C.load_asks(target)
    AA = C.ask_a(B, A)
    AB = C.ask_b(B, A)
    return B, R, AA, AB


def _c(ws, r, col, v, fmt_=None, bold=False, fill=None, align=None, italic=False, color=INK, size=10):
    c = ws.cell(row=r, column=col, value=v)
    c.font = Font(name="Calibri", size=size, bold=bold, italic=italic, color=color)
    if fmt_:
        c.number_format = fmt_
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if align:
        c.alignment = align
    return c


def _header(ws, r, labels, widths):
    for j, (lab, w) in enumerate(zip(labels, widths), 1):
        c = _c(ws, r, j, lab, bold=True, fill=HEAD,
               align=Alignment(wrap_text=True, vertical="bottom", horizontal="left" if j <= 2 else "right"))
        c.border = Border(bottom=THIN)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[r].height = 46


def write_book(path, B, R, AA, AB):
    G = R["G"]
    X = G["X"]
    pick = G["pick"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Spotlight"
    _c(ws, 1, 1, "October 2026 Sluis spotlight: %s, about %s external stars expected in its first 90 days public"
       % (pick["name"], fmt(k(pick["total"]))), bold=True, size=14)
    _c(ws, 2, 1, "Pick under the featuring memo of 15 September 2026 for the issue of Tuesday 6 October 2026, from the "
                 "warehouse extract to 27 September 2026. Expected stars to the nearest thousand.", color=GREY)
    _c(ws, 4, 1, "Spotlight", bold=True)
    _c(ws, 4, 3, "%s (%s)" % (pick["name"], pick["full"]), bold=True)
    _c(ws, 5, 1, "Expected external stars in its first 90 days public, if spotlighted")
    _c(ws, 5, 3, k(pick["total"]), "#,##0", bold=True)
    _c(ws, 6, 1, "Of which the spotlight adds")
    _c(ws, 6, 3, k(pick["lift"]), "#,##0", bold=True)
    nxt = next_names(G)
    _c(ws, 7, 1, "Next: %s, %s about %s" % (" and ".join(nxt), "each at" if len(nxt) > 1 else "at",
                                            fmt(k(G["runner"]["total"]))), color=GREY)
    hdr = ["Repository", "Owner", "Made public (UTC)", "Days public at 27 Sep", "External stars to 27 Sep",
           "Expected without the spotlight, day 90", "Stars per follower view over its life (%)",
           "Feed-card stars from day 8", "Followers of its starrers from day 8", "Stars per follower view from day 8 (%)",
           "Stars the spotlight adds", "Expected external stars, first 90 days, if spotlighted"]
    _header(ws, 9, hdr, [13, 18, 17, 11, 12, 14, 12, 12, 13, 11, 13, 16])
    r = 10
    for row in G["rows"]:
        fill = PICK if row is pick else None
        rep = B["repos"][row["rid"]]
        _c(ws, r, 1, row["name"], bold=row is pick, fill=fill)
        _c(ws, r, 2, rep["owner"]["login"], fill=fill)
        _c(ws, r, 3, B["public"][row["rid"]].strftime("%d %b %Y %H:%M"), fill=fill)
        _c(ws, r, 4, row["a"], "0", fill=fill)
        _c(ws, r, 5, int(row["S"]), "#,##0", fill=fill)
        _c(ws, r, 6, k(row["base"]), "#,##0", fill=fill)
        _c(ws, r, 7, round(100 * row["c_life"], 2), "0.00", fill=fill, color=GREY)
        _c(ws, r, 8, row["feed"], "#,##0", fill=fill)
        _c(ws, r, 9, int(row["expo"]), "#,##0", fill=fill)
        _c(ws, r, 10, round(100 * row["c"], 2), "0.00", fill=fill)
        _c(ws, r, 11, k(row["lift"]), "#,##0", fill=fill)
        _c(ws, r, 12, k(row["total"]), "#,##0", bold=True, fill=fill)
        r += 1
    r += 1
    _c(ws, r, 1, "Inputs", bold=True)
    r += 1
    inputs = [
        ("Newsletter stars per subscriber, lottery issues January 2025 to September 2026 (%)", round(100 * X["rate"], 4), "0.0000"),
        ("Subscribers on the October 2026 list", X["subs_oct"], "#,##0"),
        ("The October issue's own newsletter stars (expected)", k(X["D"]), "#,##0"),
        ("Mean followers of the accounts that starred a spotlight from the newsletter", round(X["fR"], 1), "0.0"),
        ("Mean followers of the accounts that starred from a feed card", round(X["f2"], 2), "0.00"),
    ]
    for lab, v, f in inputs:
        _c(ws, r, 1, lab)
        _c(ws, r, 10, v, f)
        r += 1
    r += 1
    spike = max(G["rows"], key=lambda q: q["c_life"] / q["c"])
    notes = [
        "Without the spotlight: external stars to 27 September plus the last seven days' daily pace times the ratio of "
        "stars still to come by day 90 to that pace, at the same age, among the lottery's runners-up.",
        "The spotlight adds the issue's own newsletter stars and the stars they set off: each star shows in the "
        "starrer's followers' home feeds, where followers star at the repository's rate per follower view, and those "
        "stars show to their followers in turn. Added stars = newsletter stars x (1 + rate x newsletter starrers' mean "
        "followers / (1 - rate x feed starrers' mean followers)).",
        "The rate per view is read from day 8 public on (feed-card stars over the followers of everyone who starred "
        "it from that day): while a repository is on an outside forum's front page, followers who see it starred have "
        "often seen it there already and star from the card more readily, and the October readers' stars will come "
        "weeks after any launch. %s's cards converted at %.2f per cent a view over its life and at %.2f per cent from "
        "day 8. On the lottery's closed spotlights the day-8 rate reproduces the stars each spotlight added to "
        "within a few hundred, where the whole-life rate misses by more than a thousand." % (
            spike["name"], 100 * spike["c_life"], 100 * spike["c"]),
    ]
    for n in notes:
        c = _c(ws, r, 1, n, color=GREY, align=Alignment(wrap_text=True, vertical="top"))
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
        ws.row_dimensions[r].height = 44
        r += 1
    ws.freeze_panes = "B10"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws2 = wb.create_sheet("Sluis issues")
    _c(ws2, 1, 1, "Sluis issues, April to September 2026: delivery and spotlight click-through", bold=True, size=13)
    _c(ws2, 2, 1, "Delivered: subscribers the issue was delivered to (the Thursday send to non-openers reaches "
                  "subscribers already counted). Clicked through: subscribers who clicked a link to the issue's "
                  "spotlighted repository themselves, by its address or its permanent address, in any of the issue's "
                  "sends; a mail gateway's check of every link of a message on delivery is not a click.", color=GREY)
    _header(ws2, 4, ["Issue", "Spotlighted repository", "Sent through", "Subscribers delivered to",
                     "Subscribers who clicked through to the spotlight", "Share who clicked through (%)"],
            [16, 30, 13, 15, 17, 15])
    r = 5
    for ym, v in AA.items():
        _c(ws2, r, 1, "%s %d" % (K.MONTH_NAME[ym[1]], ym[0]))
        _c(ws2, r, 2, v["repo"])
        _c(ws2, r, 3, K.OLD_ESP if ym[1] < 6 else K.NEW_ESP)
        _c(ws2, r, 4, v["delivered"], "#,##0")
        _c(ws2, r, 5, v["clickers"], "#,##0")
        _c(ws2, r, 6, v["share2"], "0.00")
        r += 1
    ws2.freeze_panes = "A5"

    ws3 = wb.create_sheet("Nominations")
    _c(ws3, 1, 1, "October 2026 shortlist: who nominated each repository", bold=True, size=13)
    _c(ws3, 2, 1, "People, not form responses: a signed-out response is matched to the account whose primary email "
                  "it was when it was made, a merged account to the account it was merged into, and a renamed "
                  "repository's old name to the repository. Outside: not a member of the owning organisation when "
                  "nominating, through any of their accounts, those later merged into the one they nominated from "
                  "included.", color=GREY)
    _header(ws3, 4, ["Repository", "Owner", "People who nominated it", "Of them outside its organisation",
                     "Share outside its organisation (%)"], [14, 20, 14, 15, 15])
    r = 5
    for row in G["rows"]:
        v = AB[row["name"]]
        _c(ws3, r, 1, row["name"])
        _c(ws3, r, 2, B["repos"][row["rid"]]["owner"]["login"] + ("" if v["org"] else " (user)"))
        _c(ws3, r, 3, v["people"], "#,##0")
        if v["org"]:
            _c(ws3, r, 4, v["outside"], "#,##0")
            _c(ws3, r, 5, v["share2"], "0.00")
        else:
            _c(ws3, r, 4, "owned by a user", color=GREY, align=Alignment(horizontal="right"))
        r += 1
    ws3.freeze_panes = "A5"

    ws4 = wb.create_sheet("Notes")
    ws4.column_dimensions["A"].width = 120
    notes = [
        ("Notes", True),
        ("Source: the events warehouse extract of 28 September 2026 (star events to 27 September 2026 23:59 UTC, "
         "account snapshot and organisation memberships at 27 September), the spotlight lottery log and the "
         "featuring memo of 15 September 2026.", False),
        ("External stars follow the memo's section 2: an account's first star of the repository, not a bot and not a "
         "member of the owning organisation when it stars, from any surface; day 1 is the UTC day the repository was "
         "made public.", False),
        ("Spotlight sheet: expected stars, stars without the spotlight and stars the spotlight adds are rounded to the "
         "nearest thousand from unrounded figures; stars to 27 September, feed-card stars and followers are counts.",
         False),
        ("Sluis issues: Postwerk to the May issue, Relaymark from June. Each issue went to the Sluis list in two batches "
         "on its Tuesday and again on the Thursday to subscribers who had not opened it; delivered counts the "
         "Tuesday batches, click-through counts subscribers clicking the spotlight's heading, screenshot or button, "
         "or the editor's opening paragraph's link to it, in any of the three sends; the paragraph links repositories "
         "by their permanent address on the host (/repositories/ and the id), the spotlight and in most months another "
         "repository from the shortlist. Some subscribers' mail gateways open every link of a message the second it is "
         "delivered, some of them every link but the unsubscribe link: Relaymark logs that as four or more links "
         "clicked within two seconds, Postwerk as one click on every link of the send or on every link but the "
         "unsubscribe link, and those fetches are left out (the platforms' own click counts leave out only "
         "the gateways they recognise). Test sends to the editorial seed addresses and Product News are left out.",
         False),
        ("Nominations: responses to the October form. A signed-out response is matched to the account whose primary "
         "email it was when it was made, as the account snapshot or a signed-in response of that account records it, "
         "a merged account to the account it was merged into, kiln-ui to kilnkit and imagelint to hullcheck; "
         "addresses of forks under other owners are not nominations of the shortlisted repository. Five nominators "
         "used a personal account while still members through a work account that they merged into it when they left, "
         "a few days later: they nominated from inside. Five more nominated their organisation's repository signed in "
         "and again signed out with their work address, which their account no longer carries since they left: one "
         "person each, inside.", False),
    ]
    for i, (text, bold) in enumerate(notes, 1):
        _c(ws4, i, 1, text, bold=bold, align=Alignment(wrap_text=True, vertical="top"))
        ws4.row_dimensions[i].height = 16 if bold else 15 * (1 + len(text) // 135)
    wb.save(path)
    PP.normalize_ooxml(path, K.PEOPLE["editor"], K.PEOPLE["editor"], "2026-09-28T07:05:12Z", "2026-09-29T15:22:40Z",
                       excel=True, company=K.ORG)


def write_chart(path, B, R):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    G = R["G"]
    rows = G["rows"]
    pick = G["pick"]
    names = [r["name"] for r in rows]
    organic = [k(r["base"]) for r in rows]
    added = [k(r["lift"]) for r in rows]
    exp = [k(r["total"]) for r in rows]
    now = [int(r["S"]) for r in rows]
    fig, ax = plt.subplots(figsize=(9.6, 5.6), dpi=110)
    y = list(range(len(rows)))[::-1]
    h = 0.38
    ax.barh([v + h / 2 for v in y], organic, height=h, color="#5b8c99", label="Expected by day 90 without the spotlight")
    ax.barh([v + h / 2 for v in y], added, left=organic, height=h,
            color=["#d4872c" if r is pick else "#9fc3cc" for r in rows], label="Added by the October spotlight")
    ax.barh([v - h / 2 for v in y], now, height=h * 0.8, color="#c9d1d5", label="External stars by 27 September 2026")
    for v, e, n, r in zip(y, exp, now, rows):
        ax.text(e + 140, v + h / 2, fmt(e), va="center", fontsize=9.5, color="#1d2733",
                fontweight="bold" if r is pick else "normal")
        ax.text(n + 140, v - h / 2, fmt(k(n)), va="center", fontsize=8.5, color="#5a6470")
    ax.set_yticks(y)
    ax.set_yticklabels(names, fontsize=10)
    for lab in ax.get_yticklabels():
        if lab.get_text() == pick["name"]:
            lab.set_fontweight("bold")
    ax.set_xlim(0, max(exp) * 1.13)
    ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, p: fmt(v)))
    ax.set_xlabel("External stars in the first 90 days public (to the nearest thousand)", fontsize=9)
    ax.set_title("%s's spotlight would add about %s stars and put it first at about %s"
                 % (pick["name"], fmt(k(pick["lift"])), fmt(k(pick["total"]))), fontsize=11.5, loc="left")
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(loc="lower right", fontsize=8.3, frameon=False)
    fig.text(0.01, 0.012, "Sluis October 2026 shortlist, extract to 27 September 2026. Expected stars if spotlighted: "
                          "stars to date, the organic forecast to day 90 and the spotlight's newsletter and feed stars.",
             fontsize=7.6, color="#5a6470")
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    fig.savefig(path, metadata={"Software": None, "Title": None, "Author": None, "Description": None,
                                "Creation Time": None})
    plt.close(fig)


def figures(B, R, AA, AB):
    G = R["G"]
    X = G["X"]
    D = R["out"]
    w0 = R["ranks"]["R0"]["winner"]
    w3 = R["ranks"]["R3"]["winner"]
    f = {
        "call": {"pick": G["pick"]["name"], "pick_full_name": G["pick"]["full"],
                 "expected_external_stars_90d": k(G["pick"]["total"]),
                 "stars_the_spotlight_adds": k(G["pick"]["lift"]),
                 "next": next_names(G), "next_expected_stars": k(G["runner"]["total"])},
        "candidates": {r["name"]: {"expected_external_stars_90d": k(r["total"]),
                                   "expected_without_spotlight": k(r["base"]),
                                   "stars_the_spotlight_adds": k(r["lift"]),
                                   "external_stars_to_2026_09_27": int(r["S"]),
                                   "external_stars_to_2026_09_27_thousands": k(r["S"]),
                                   "days_public_at_cut": r["a"],
                                   "feed_card_stars_from_day_8": r["feed"],
                                   "followers_reached_from_day_8": int(r["expo"]),
                                   "stars_per_follower_view_from_day_8_pct": round(100 * r["c"], 2),
                                   "stars_per_follower_view_whole_life_pct": round(100 * r["c_life"], 2)}
                       for r in G["rows"]},
        "conventions": {"window_days": K.WINDOW_DAYS, "pace_window_days": X["win"], "lottery_issues": len(X["past"]),
                        "rate_from_day": C.POST_FROM},
        "inputs": {"newsletter_stars_per_subscriber_pct": round(100 * X["rate"], 4),
                   "october_subscribers": X["subs_oct"],
                   "october_direct_newsletter_stars": k(X["D"]),
                   "newsletter_starrers_mean_followers": round(X["fR"], 1),
                   "feed_starrers_mean_followers": round(X["f2"], 2)},
        "backtest": {"spotlights": len(R["L"]), "rmse_rate_from_day_8": int(round(R["backtest"]["post"], -1)),
                     "rmse_rate_whole_life": int(round(R["backtest"]["life"], -1))},
        "default_path": {"pick": w0, "expected_external_stars_90d": k(D[w0]["R0"]),
                         "stars_the_spotlight_adds": k(R["Lbar"])},
        "whole_life_rate": {"pick": w3, "expected_external_stars_90d": k(D[w3]["R3"])},
        "ask_a": {"%d-%02d" % ym: {"delivered": v["delivered"], "clicked_through": v["clickers"],
                                   "share_pct": v["share2"]} for ym, v in AA.items()},
        "ask_b": {name: ({"people": v["people"], "outside": v["outside"], "share_outside_pct": v["share2"]}
                         if v["org"] else {"people": v["people"]}) for name, v in AB.items()},
    }
    return f


def main(target, golden_dir, figures_path):
    os.makedirs(golden_dir, exist_ok=True)
    B, R, AA, AB = compute(target)
    write_book(os.path.join(golden_dir, K.GOLD_BOOK), B, R, AA, AB)
    write_chart(os.path.join(golden_dir, K.GOLD_CHART), B, R)
    f = figures(B, R, AA, AB)
    with open(figures_path, "w") as fh:
        json.dump(f, fh, indent=2)
        fh.write("\n")
    stamp = K.PACK_STAMP.timestamp() + 2 * 86400
    for fn in os.listdir(golden_dir):
        os.utime(os.path.join(golden_dir, fn), (stamp, stamp))
    return f
