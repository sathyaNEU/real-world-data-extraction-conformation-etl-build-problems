"""The pack's documents: the featuring memo (PDF) and the spotlight lottery log (workbook)."""
from __future__ import annotations

import datetime as dt

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import A4                                                # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import cm                                                    # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import common as K   # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/"
for _n, _f in (("Carlito", "crosextra/Carlito-Regular.ttf"), ("Carlito-Bold", "crosextra/Carlito-Bold.ttf"),
               ("Carlito-Italic", "crosextra/Carlito-Italic.ttf")):
    pdfmetrics.registerFont(TTFont(_n, FONT_DIR + _f))
SANS, SANS_B, SANS_I = "Carlito", "Carlito-Bold", "Carlito-Italic"
INK = colors.HexColor("#1d2733")
GREY = colors.HexColor("#5a6470")
TEAL = colors.HexColor("#1f5f6b")
RULE = colors.HexColor("#b4bec6")
PALE = colors.HexColor("#e9eff2")

BODY = ParagraphStyle("body", fontName=SANS, fontSize=10.6, leading=14.2, textColor=INK, spaceAfter=6)
CLAUSE = ParagraphStyle("clause", parent=BODY, leftIndent=28, firstLineIndent=-28)
H1 = ParagraphStyle("h1", fontName=SANS_B, fontSize=16, leading=20, textColor=TEAL, spaceAfter=4)
SUB = ParagraphStyle("sub", fontName=SANS, fontSize=10, leading=13, textColor=GREY, spaceAfter=10)
H2 = ParagraphStyle("h2", fontName=SANS_B, fontSize=11.6, leading=15, textColor=INK, spaceBefore=10, spaceAfter=4)
CELL = ParagraphStyle("cell", fontName=SANS, fontSize=9.6, leading=12, textColor=INK)
CELLB = ParagraphStyle("cellb", fontName=SANS_B, fontSize=9.6, leading=12, textColor=INK)


def _footer(canvas, doc):
    canvas.saveState()
    canvas.setFont(SANS, 8)
    canvas.setFillColor(GREY)
    canvas.drawString(2.2 * cm, 1.25 * cm, "Grachtwerk B.V.  |  Sluis editorial  |  internal")
    canvas.drawRightString(A4[0] - 2.2 * cm, 1.25 * cm, "Featuring memo, 15 September 2026  |  page %d" % doc.page)
    canvas.restoreState()


def write_memo(path):
    P = K.PEOPLE
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=2.3 * cm, rightMargin=2.3 * cm, topMargin=2.0 * cm,
                            bottomMargin=2.1 * cm, title="Sluis featuring memo", author=P["editor"], subject="",
                            creator="", producer="")
    s = [Paragraph("Sluis: choosing the spotlight from October 2026", H1),
         Paragraph("Featuring memo, 15 September 2026", SUB)]
    head = [["From", "%s, editor of Sluis" % P["editor"]],
            ["To", "Sluis editorial team"],
            ["Copy", "%s (developer relations), %s (growth), %s (data engineering)" % (P["devrel"], P["growth"], P["data"])],
            ["Subject", "How each issue's spotlighted repository is chosen"]]
    t = Table([[Paragraph(a, CELLB), Paragraph(b, CELL)] for a, b in head], colWidths=[2.4 * cm, 13.6 * cm], hAlign="LEFT")
    t.setStyle(TableStyle([("LINEBELOW", (0, -1), (-1, -1), 0.6, RULE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    s += [t, Spacer(1, 8)]

    def cl(no, text):
        s.append(Paragraph("%s&nbsp;&nbsp;&nbsp;%s" % (no, text), CLAUSE))

    s.append(Paragraph("1. Scope", H2))
    cl("1.1", "This memo sets how the repository spotlighted in each issue of Sluis is chosen, from the issue of "
              "Tuesday 6 October 2026 onwards. It replaces the spotlight lottery the growth team has run since the "
              "January 2025 issue; the September 2026 issue was the last one drawn.")
    cl("1.2", "The rest of the spotlight stays as it is. The editors shortlist seven repositories from the month's "
              "reader nominations, the spotlighted repository gets the issue's write-up and link, and its maintainers "
              "receive a year of %s credit from developer relations." % K.CREDIT)
    s.append(Paragraph("2. Which repository is spotlighted", H2))
    cl("2.1", "Each issue's spotlight goes to the shortlisted repository with the most expected external stars in its "
              "first 90 days public if it is that issue's spotlight.")
    cl("2.2", "An external star is an account's first star of the repository, made by an account that is neither a bot "
              "account nor a member of the organisation that owns the repository at the time it stars. A star counts "
              "whichever surface it is made from. An account that removes its star and stars the repository again has "
              "made one external star.")
    cl("2.3", "A repository's first 90 days public start on the UTC calendar day on which it was made public, which is "
              "day 1, and end at the end of day 90.")
    cl("2.4", "The expected external stars count every external star the repository will have by the end of its 90th "
              "day, including those it has not yet received when the pick is made.")
    s.append(Paragraph("3. The October 2026 issue", H2))
    cl("3.1", "Nominations for the October issue close on Sunday 20 September 2026 and the shortlist of seven on "
              "Wednesday 23 September.")
    cl("3.2", "The pick is made from the events warehouse extract to Sunday 27 September 2026 and locked at the "
              "editorial meeting on Wednesday 30 September.")
    cl("3.3", "The issue goes out at 09:00 Amsterdam time on Tuesday 6 October 2026.")
    s.append(Paragraph("4. Later issues", H2))
    cl("4.1", "Later issues keep the same timetable, counted back from the first Tuesday of the month.")
    s.append(Paragraph("5. Records", H2))
    cl("5.1", "The lottery log stays with the growth team and the star history in the events warehouse.")
    cl("5.2", "The editor records each pick in the lottery log's sheet for the year.")
    s.append(Spacer(1, 14))
    s.append(Paragraph("%s<br/>Editor, Sluis" % P["editor"], BODY))
    doc.build(s, onFirstPage=_footer, onLaterPages=_footer)


# ---------------------------------------------------------------------------------------------------- lottery log

def shortlist_close(ym):
    """The Wednesday thirteen days before the issue goes out."""
    return K.first_tuesday(*ym) - dt.timedelta(days=13)


def lottery_rows(W):
    by_issue = {}
    for r in W.lottery:
        by_issue.setdefault(r.issue, []).append(r)
    rows = []
    for ym in K.LOTTERY_MONTHS + [K.OCT]:
        reps = sorted(W.october if ym == K.OCT else by_issue[ym], key=lambda r: K.hexid("order", r.id))
        drawn = None if ym == K.OCT else [r for r in reps if r.drawn][0]
        g = K.rng("draw-time", *ym)
        draw_at = dt.datetime.combine(shortlist_close(ym) + dt.timedelta(days=1), dt.time(10, 0)) + \
            dt.timedelta(minutes=int(g.integers(0, 300)))
        rows.append(dict(ym=ym, reps=reps, drawn=drawn, draw_at=draw_at, subs=W.list[ym]))
    return rows


def write_lottery(path, W):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2025"
    bold = Font(name="Calibri", size=11, bold=True)
    ws["A1"] = "Sluis spotlight lottery 2025"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True)
    ws["A2"] = "Seven shortlisted per issue; one drawn by lot (a number from 1 to 7, in shortlist order)."
    hdr = ["Issue", "Shortlist closed", "Subscribers"] + ["Shortlist %d" % i for i in range(1, 8)] + \
        ["Drawn", "Draw no.", "Drawn at"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = bold
        c.fill = PatternFill("solid", fgColor="DDE7EA")
    rows = lottery_rows(W)
    i = 5
    for x in rows:
        if x["ym"][0] != 2025:
            continue
        ws.cell(row=i, column=1, value=K.first_tuesday(*x["ym"])).number_format = "d mmm yyyy"
        ws.cell(row=i, column=2, value=shortlist_close(x["ym"])).number_format = "d mmm yyyy"
        ws.cell(row=i, column=3, value=x["subs"]).number_format = "#,##0"
        for j, r in enumerate(x["reps"]):
            ws.cell(row=i, column=4 + j, value="%s/%s" % (r.owner, r.name))
        ws.cell(row=i, column=11, value="%s/%s" % (x["drawn"].owner, x["drawn"].name))
        ws.cell(row=i, column=12, value=x["reps"].index(x["drawn"]) + 1)
        ws.cell(row=i, column=13, value=x["draw_at"]).number_format = "d mmm yyyy hh:mm"
        i += 1
    for j, w in enumerate([12, 15, 12] + [27] * 7 + [27, 9, 17], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "D5"

    ws2 = wb.create_sheet("2026")
    hdr2 = ["issue_date", "shortlist_closed", "subscribers"] + ["repo_id_%d" % i for i in range(1, 8)] + \
        ["drawn_repo_id", "draw_no", "drawn_at", "note"]
    for j, h in enumerate(hdr2, 1):
        c = ws2.cell(row=1, column=j, value=h)
        c.font = bold
    i = 2
    for x in rows:
        if x["ym"][0] != 2026:
            continue
        ws2.cell(row=i, column=1, value=K.first_tuesday(*x["ym"])).number_format = "yyyy-mm-dd"
        ws2.cell(row=i, column=2, value=shortlist_close(x["ym"])).number_format = "yyyy-mm-dd"
        ws2.cell(row=i, column=3, value=x["subs"])
        for j, r in enumerate(x["reps"]):
            ws2.cell(row=i, column=4 + j, value=r.pub)
        if x["drawn"] is not None:
            ws2.cell(row=i, column=11, value=x["drawn"].pub)
            ws2.cell(row=i, column=12, value=x["reps"].index(x["drawn"]) + 1)
            ws2.cell(row=i, column=13, value=x["draw_at"]).number_format = "yyyy-mm-dd hh:mm"
        else:
            ws2.cell(row=i, column=14, value="no draw: the October pick follows the featuring memo of 15 Sep 2026 "
                                              "(extract to 27 Sep, editorial meeting 30 Sep)")
        i += 1
    for j, w in enumerate([12, 16, 12] + [13] * 7 + [15, 9, 17, 70], 1):
        ws2.column_dimensions[get_column_letter(j)].width = w
    ws2.freeze_panes = "D2"
    wb.save(path)
