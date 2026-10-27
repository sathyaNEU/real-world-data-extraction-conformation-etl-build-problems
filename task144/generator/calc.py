"""Every figure from the shipped pack alone: the October pick and its rivals (calc.golden, calc.rungs), the
default path, and the two asks. Reads only files under target/."""
from __future__ import annotations

import datetime as dt
import json
import math
import os
import re

import numpy as np
import openpyxl
import pandas as pd

import common as K

OCT_ISSUE = dt.date(2026, 10, 6)


# ---------------------------------------------------------------------------------------------------- loading

def load(target):
    p = lambda f: os.path.join(target, f)
    B = {}
    xw = pd.read_csv(p(K.CROSSWALK))
    code = dict(zip(xw.v5_src, xw.v6_surface))
    e5 = pd.read_csv(p(K.EV25))
    e5 = pd.DataFrame({"t": pd.to_datetime(e5.event_ts_ms, unit="ms"), "repo": e5.repo_id, "account": e5.actor_id,
                       "action": e5.action, "surface": e5.src.map(code), "seq": np.arange(len(e5))})
    e6 = pd.read_csv(p(K.EV26))
    e6 = pd.DataFrame({"t": pd.to_datetime(e6.occurred_at.str.rstrip("Z")), "repo": e6.repository_id,
                       "account": e6.account_id, "action": e6.event_type.map({"star.created": "star",
                                                                              "star.deleted": "unstar"}),
                       "surface": e6.surface, "seq": np.arange(len(e6)) + len(e5)})
    ev = pd.concat([e5, e6], ignore_index=True).sort_values(["t", "seq"], kind="mergesort").reset_index(drop=True)
    B["events"] = ev
    B["unmapped_codes"] = int(ev.surface.isna().sum())
    B["accounts"] = pd.read_parquet(p(K.ACCOUNTS))
    acc = B["accounts"].set_index("account_id")
    B["followers"] = acc.followers.to_dict()
    B["acc_type"] = acc.account_type.to_dict()
    mem = pd.read_csv(p(K.MEMBERS), keep_default_na=False)
    mem["joined"] = pd.to_datetime(mem.joined_at.str.rstrip("Z"))
    mem["left"] = pd.to_datetime(mem.left_at.replace("", None).str.rstrip("Z"))
    B["members"] = mem
    repos = json.load(open(p(K.REPOS), encoding="utf-8"))
    B["repos"] = {r["id"]: r for r in repos}
    B["public"] = {r["id"]: dt.datetime.strptime(r["made_public_at"], "%Y-%m-%dT%H:%M:%SZ") for r in repos}
    roll = pd.read_csv(p(K.ROLLUP))
    roll["day"] = pd.to_datetime(roll.day_utc)
    B["rollup"] = roll
    B["lottery"] = load_lottery(p(K.LOTTERY), B)
    B["cut"] = K.CUT
    return B


def load_lottery(path, B):
    wb = openpyxl.load_workbook(path, data_only=True)
    by_full = {r["full_name"]: rid for rid, r in B["repos"].items()}
    out = []
    ws = wb["2025"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [i for i, r in enumerate(rows) if r and r[0] == "Issue"][0]
    for r in rows[hdr + 1:]:
        if not r or r[0] is None:
            continue
        out.append(dict(issue=r[0].date() if isinstance(r[0], dt.datetime) else r[0], subs=int(r[2]),
                        shortlist=[by_full[x] for x in r[3:10]], drawn=by_full[r[10]]))
    ws = wb["2026"]
    rows = list(ws.iter_rows(values_only=True))
    for r in rows[1:]:
        if not r or r[0] is None:
            continue
        out.append(dict(issue=r[0].date() if isinstance(r[0], dt.datetime) else r[0], subs=int(r[2]),
                        shortlist=[int(x) for x in r[3:10]], drawn=int(r[10]) if r[10] is not None else None))
    return out


# ---------------------------------------------------------------------------------------------------- stars

def send_utc(issue_date):
    """09:00 Amsterdam time on the issue's Tuesday, in UTC."""
    return K.from_amsterdam(dt.datetime.combine(issue_date, dt.time(9, 0)))


def owner_of(B, rid):
    o = B["repos"][rid]["owner"]
    return o["login"] if o["type"] == "Organization" else None


def first_stars(B, opts=None):
    """Each account's first star of each repository (the earliest star event of the pair)."""
    opts = opts or {}
    ev = B["events"]
    st = ev[ev.action == "star"]
    if opts.get("all_events"):
        f = st.copy()
    else:
        f = st.drop_duplicates(["repo", "account"], keep="first").copy()
    if opts.get("net_unstars"):
        un = ev[ev.action == "unstar"][["repo", "account"]].drop_duplicates()
        last = ev.drop_duplicates(["repo", "account"], keep="last")
        gone = set(map(tuple, last[last.action == "unstar"][["repo", "account"]].values))
        f = f[[(r, a) not in gone for r, a in zip(f.repo, f.account)]]
    return f


def is_member(B, org, account, t, current=False):
    m = B["members_by"].get((org, account))
    if not m:
        return False
    for joined, left in m:
        if current:
            if pd.isna(left):
                return True
        elif joined <= t and (pd.isna(left) or t < left):
            return True
    return False


def prepare(B):
    if "members_by" not in B:
        mb = {}
        for org, aid, j, l in zip(B["members"].organization, B["members"].account_id, B["members"].joined,
                                  B["members"].left):
            mb.setdefault((org, aid), []).append((j, l))
        B["members_by"] = mb


EXT_KEYS = ("all_events", "net_unstars", "keep_bots", "keep_members", "members_now", "hour_window", "amsterdam_days",
            "stars_to_survivor")


def external(B, opts=None):
    """The memo's external stars: first stars from accounts that are not bots and not members of the owning
    organisation when they star. opts switch the readings the fork table lists."""
    opts = opts or {}
    key = tuple(bool(opts.get(k)) for k in EXT_KEYS)
    cache = B.setdefault("_ext", {})
    if key not in cache:
        cache[key] = _external(B, opts)
    return cache[key]


def _external(B, opts):
    prepare(B)
    f = first_stars(B, opts)
    types = B["acc_type"]
    keep = []
    accs = B["accounts"]
    survivor = dict(zip(accs.account_id[accs.merged_into.notna()], accs.merged_into.dropna().astype(int)))
    for rid, aid, t in zip(f.repo, f.account, f.t):
        if opts.get("stars_to_survivor"):
            # the reading that credits a merged account's star to the account it was later merged into
            while aid in survivor:
                aid = survivor[aid]
        if not opts.get("keep_bots") and types.get(aid) == "bot":
            keep.append(False)
            continue
        org = owner_of(B, rid)
        if org and not opts.get("keep_members") and is_member(B, org, aid, t, current=opts.get("members_now", False)):
            keep.append(False)
            continue
        keep.append(True)
    f = f[keep].copy()
    f["day"] = [day_index(B["public"][r], t, opts) for r, t in zip(f.repo, f.t)]
    return f


def day_index(public, t, opts=None):
    if opts and opts.get("hour_window"):
        return int((t - public).total_seconds() // 86400) + 1
    if opts and opts.get("amsterdam_days"):
        return (K.to_amsterdam(t).date() - K.to_amsterdam(public).date()).days + 1
    return (t.date() - public.date()).days + 1


def daily_from_events(B, ext, rid, last):
    d = np.zeros(102)
    sub = ext[(ext.repo == rid) & (ext.t <= B["cut"])]
    for k in sub.day:
        if 1 <= k <= last:
            d[k] += 1
    return d


def daily_from_rollup(B, rid):
    d = np.zeros(102)
    sub = B["rollup"][B["rollup"].repository_id == rid]
    pub = B["public"][rid].date()
    for day, n in zip(sub.day, sub.external_first_stars):
        k = (day.date() - pub).days + 1
        if 1 <= k <= 100:
            d[k] = n
    return d


# ---------------------------------------------------------------------------------------------------- the pick

def context(B, opts=None):
    """Everything the constructions share: the candidates, the runners-up's future-to-pace ratio, the lottery's
    direct newsletter rate and the reach of newsletter and feed starrers."""
    opts = opts or {}
    lot = B["lottery"]
    octrow = [x for x in lot if x["issue"] == OCT_ISSUE][0]
    past = [x for x in lot if x["drawn"] is not None]
    cands = octrow["shortlist"]
    cut_day = B["cut"].date()
    runners = [rid for x in past for rid in x["shortlist"] if rid != x["drawn"]]
    complete = [rid for rid in runners if B["public"][rid].date() + dt.timedelta(days=89) <= cut_day]
    dd = {rid: daily_from_rollup(B, rid) for rid in complete}
    win = opts.get("pace_window", 7)

    def G(a):
        num = sum(dd[r][a + 1:91].sum() for r in complete)
        den = sum(dd[r][a - win + 1:a + 1].sum() / win for r in complete)
        return num / den

    ext = external(B, opts)
    fl = B["followers"]
    spots = [x for x in past]
    issues_for_rate = spots
    nl_found = set(ext[ext.surface == "newsletter"].repo)
    if opts.get("rate_issues") == "found":
        issues_for_rate = [x for x in spots if x["drawn"] in nl_found]
    if opts.get("rate_issues") == "v6":
        issues_for_rate = [x for x in spots if x["issue"] >= K.SCHEMA_V6.date()]
    if opts.get("rate_issues") == "2026":
        issues_for_rate = [x for x in spots if x["issue"].year == 2026]
    elif opts.get("rate_issues") == "last6":
        issues_for_rate = spots[-6:]
    nl = ext[ext.surface == "newsletter"]
    nl_by = nl.groupby("repo").size().to_dict()
    rate = sum(nl_by.get(x["drawn"], 0) for x in issues_for_rate) / sum(x["subs"] for x in issues_for_rate)
    subs_oct = octrow["subs"] if not opts.get("list_september") else [x for x in lot if x["issue"] == dt.date(2026, 9, 1)][0]["subs"]
    D = rate * subs_oct
    fs = first_stars(B)
    spot_ids = {x["drawn"] for x in spots}
    nlf = fs[(fs.surface == "newsletter") & fs.repo.isin(spot_ids)]
    if opts.get("fR") == "2026":
        nlf = nlf[nlf.repo.isin({x["drawn"] for x in spots if x["issue"].year == 2026})]
    if opts.get("fR") == "issue_mean":
        fR = float(np.mean([np.mean([fl[a] for a in g.account]) for _, g in nlf.groupby("repo")]))
    else:
        fR = float(np.mean([fl[a] for a in nlf.account]))
    fd = fs[fs.surface == "feed"]
    if opts.get("f2") == "candidates":
        fd = fd[fd.repo.isin(set(cands))]
    elif opts.get("f2") == "spotlights":
        fd = fd[fd.repo.isin(spot_ids)]
    f2 = float(np.mean([fl[a] for a in fd.account]))
    return dict(cands=cands, octrow=octrow, past=past, complete=complete, dd=dd, G=G, ext=ext, rate=rate, D=D,
                fR=fR, f2=f2, subs_oct=subs_oct, fs=fs, win=win)


POST_FROM = 8     # the rate per view after a repository's first week: its stars from its 8th day public


def view_rate(B, fs, rid, until, opts=None, life=False):
    """A repository's rate per follower view: its feed-card stars over the followers of every account that starred
    it (each account's first star), over the stars made before `until`. By default only the stars from its 8th day
    public count (life=True takes its whole life); opts move the window for the readings the checks spread over
    (post_from, last_days, lag, tail)."""
    opts = opts or {}
    fl = B["followers"]
    pub = B["public"][rid]
    mine = fs[(fs.repo == rid) & (fs.t < until)]
    if life:
        feed = int((mine.surface == "feed").sum())
        expo = float(sum(fl[q] for q in mine.account))
        return feed / expo, feed, expo
    day = np.array([(t.date() - pub.date()).days + 1 for t in mine.t])
    a0 = opts.get("post_from", POST_FROM)
    if opts.get("last_days"):
        a0 = (until - dt.timedelta(seconds=1)).date().toordinal() - pub.date().toordinal() + 2 - opts["last_days"]
    lag, tail = opts.get("lag", 0), opts.get("tail", 0)
    is_feed = (mine.surface == "feed").values
    feed = int((is_feed & (day >= a0 + lag)).sum())
    keep = (day >= a0) & (mine.t < until - dt.timedelta(days=tail)).values
    expo = float(sum(fl[q] for q in mine.account[keep]))
    return feed / expo, feed, expo


def mult(X, c, reach=None):
    """Stars the spotlight sets off per newsletter star: the reader's own, then the cascade through its followers'
    feeds at the rate per view c, every feed star shown to its own followers in turn."""
    return 1 + c * (X["fR"] if reach is None else reach) / (1 - c * X["f2"])


def candidate(B, X, rid, opts=None):
    opts = opts or {}
    pub = B["public"][rid]
    a = day_index(pub, B["cut"], opts)
    d = daily_from_events(B, X["ext"], rid, a)
    win = X["win"]
    S = d[1:a + 1].sum()
    pace = d[a - win + 1:a + 1].mean()
    base = S + pace * X["G"](a)
    until = B["cut"] + dt.timedelta(seconds=1)
    c, feed, expo = view_rate(B, X["ext"] if opts.get("rate_external") else X["fs"], rid, until, opts)
    c_life, feed_life, expo_life = view_rate(B, X["fs"], rid, until, life=True)
    ext_m = X["ext"][(X["ext"].repo == rid) & (X["ext"].t <= B["cut"])]
    Kf = float((ext_m.surface == "feed").mean())
    m = mult(X, c)
    lift = X["D"] * m
    return dict(rid=rid, name=B["repos"][rid]["name"], full=B["repos"][rid]["full_name"], a=a, S=S, pace=pace,
                base=base, feed=feed, expo=expo, c=c, c_life=c_life, feed_life=feed_life, expo_life=expo_life,
                K=Kf, m=m, lift=lift, total=base + lift, daily=d)


def golden(B, opts=None):
    X = context(B, opts)
    rows = [candidate(B, X, rid, opts) for rid in X["cands"]]
    rows.sort(key=lambda r: -r["total"])
    return dict(X=X, rows=rows, pick=rows[0], runner=rows[1], margin=rows[0]["total"] / rows[1]["total"])


def lottery_lifts(B, X, opts=None):
    """Each spotlighted repository with a complete 90-day window: its 90-day external stars against its own forecast
    from its trajectory before the issue, beside the rates per view it showed before the issue."""
    out = []
    for x in X["past"]:
        rid = x["drawn"]
        pub = B["public"][rid]
        if pub.date() + dt.timedelta(days=89) > B["cut"].date():
            continue
        send = send_utc(x["issue"])
        a = day_index(pub, send) - 1
        d = daily_from_rollup(B, rid)
        win = X["win"]
        pace = d[a - win + 1:a + 1].mean()
        base = d[1:a + 1].sum() + pace * X["G"](a)
        S90 = d[1:91].sum()
        ext = X["ext"]
        pre = ext[(ext.repo == rid) & (ext.t < send)]
        c_life = view_rate(B, X["fs"], rid, send, life=True)[0]
        c_post = view_rate(B, X["fs"], rid, send, opts)[0]
        Dk = X["rate"] * x["subs"]
        out.append(dict(rid=rid, name=B["repos"][rid]["name"], issue=x["issue"], subs=x["subs"], a=a,
                        S=d[1:a + 1].sum(), pace=pace, base=base, S90=S90, lift=S90 - base,
                        K=float((pre.surface == "feed").mean()), c_life=float(c_life), c_post=float(c_post),
                        D=Dk, pred_life=Dk * mult(X, c_life), pred_post=Dk * mult(X, c_post),
                        owner_f=owner_followers(B, rid)))
    return out


def owner_followers(B, rid):
    login = B["repos"][rid]["owner"]["login"]
    acc = B["accounts"]
    hit = acc.followers[acc.login == login]
    return float(hit.iloc[0]) if len(hit) else 0.0


RUNGS = ["R0", "R1", "R2", "R3", "R3k", "R4", "pace", "direct", "first", "platform", "analogs", "regression", "owner"]


def rungs(B, opts=None):
    """The default path, its corrections and the rivals, each a full construction from the pack. R4 is the golden;
    R3 is the same cascade on each repository's rate per view over its whole life, R3k that cascade scaled to the
    lottery's lifts."""
    G0 = golden(B, opts)
    X = G0["X"]
    L = lottery_lifts(B, X, opts)
    scale = X["subs_oct"] / np.mean([l["subs"] for l in L])
    Lbar = float(np.mean([l["lift"] for l in L]) * scale)
    pbar = float(sum(l["lift"] for l in L) / sum(l["base"] for l in L))
    rbar = float(sum(l["lift"] for l in L) / sum(l["pace"] for l in L) * scale)
    users = B["accounts"][B["accounts"].account_type == "user"]
    Fplat = float(users.followers.mean())
    obs = np.array([l["lift"] for l in L])
    cas = np.array([l["pred_life"] - l["D"] for l in L])
    kappa = float(((obs - np.array([l["D"] for l in L])) * cas).sum() / (cas * cas).sum())
    rmse = lambda p: float(np.sqrt(np.mean((obs - p) ** 2)))
    backtest = dict(life=rmse(np.array([l["pred_life"] for l in L])), post=rmse(np.array([l["pred_post"] for l in L])),
                    mean=rmse(np.full(len(obs), obs.mean())),
                    corr_life=float(np.corrcoef(obs, [l["pred_life"] for l in L])[0, 1]),
                    corr_post=float(np.corrcoef(obs, [l["pred_post"] for l in L])[0, 1]))
    # analogs: the three spotlights nearest on stars, pace and feed share before the issue
    feats = np.array([[l["S"], l["pace"], l["K"]] for l in L])
    mu, sd = feats.mean(0), feats.std(0)
    lifts_scaled = np.array([l["lift"] * X["subs_oct"] / l["subs"] for l in L])
    Xr = np.column_stack([np.ones(len(L)), feats])
    beta = np.linalg.lstsq(Xr, lifts_scaled, rcond=None)[0]
    # the owner's audience: each spotlight's lift on the log of its owner's followers
    Xo = np.column_stack([np.ones(len(L)), np.log1p([l["owner_f"] for l in L])])
    beta_o = np.linalg.lstsq(Xo, lifts_scaled, rcond=None)[0]
    out = {}
    for r in G0["rows"]:
        D, f2 = X["D"], X["f2"]
        z = (np.array([r["S"], r["pace"], r["K"]]) - mu) / sd
        dist = np.sqrt((((feats - mu) / sd - z) ** 2).sum(1))
        near = np.argsort(dist)[:3]
        out[r["name"]] = dict(
            R0=r["base"] + Lbar,
            R1=r["base"] * (1 + pbar),
            R2=r["base"] + D / (1 - r["K"]),
            R3=r["base"] + D * mult(X, r["c_life"]),
            R3k=r["base"] + D * (1 + kappa * (mult(X, r["c_life"]) - 1)),
            R4=r["total"],
            pace=r["base"] + rbar * r["pace"],
            direct=r["base"] + D,
            first=r["base"] + D * (1 + r["c"] * X["fR"]),
            platform=r["base"] + D * mult(X, r["c"], reach=Fplat),
            analogs=r["base"] + float(lifts_scaled[near].mean()),
            regression=r["base"] + float(beta @ np.array([1, r["S"], r["pace"], r["K"]])),
            owner=r["base"] + float(beta_o @ np.array([1, np.log1p(owner_followers(B, r["rid"]))])),
            base=r["base"])
    ranks = {}
    pick = G0["pick"]["name"]
    for rung in RUNGS:
        order = sorted(out, key=lambda k: -out[k][rung])
        ranks[rung] = dict(order=order, winner=order[0], margin=out[order[0]][rung] / out[order[1]][rung],
                           golden_place=order.index(pick) + 1)
    return dict(G=G0, L=L, Lbar=Lbar, pbar=pbar, rbar=rbar, Fplat=Fplat, beta=beta, beta_o=beta_o, kappa=kappa,
                backtest=backtest, out=out, ranks=ranks)


def default_path(B):
    R = rungs(B)
    w = R["ranks"]["R0"]["winner"]
    return dict(pick=w, total=R["out"][w]["R0"], added=R["Lbar"], R=R)


# ---------------------------------------------------------------------------------------------------- Ask A

ASK_MONTHS = [(2026, m) for m in range(4, 10)]


def load_asks(target):
    p = lambda f: os.path.join(target, f)
    A = {}
    A["pw_camp"] = pd.read_csv(p(K.PW_CAMP), sep=";", encoding="utf-8-sig", keep_default_na=False)
    A["pw_links"] = pd.read_csv(p(K.PW_LINKS), sep=";", encoding="utf-8-sig")
    rm = json.load(open(p(K.RM_CAMP), encoding="utf-8"))
    A["rm_camp"] = pd.DataFrame([dict(campaign_id=c["id"], campaign_name=c["name"], list=c["audience"]["list"],
                                      segment=c["audience"]["segment"] or "", sent_at=c["sent_at"],
                                      recipients=c["stats"]["recipients"], delivered=c["stats"]["delivered"])
                                 for c in rm["campaigns"]])
    A["rm_links"] = pd.read_csv(p(K.RM_LINKS))
    A["noms"] = pd.read_csv(p(K.NOMS), keep_default_na=False, dtype=str)
    A["noms"]["account_id"] = pd.Series([int(v) if v else None for v in A["noms"].account_id], dtype=object)
    A["renames"] = pd.read_csv(p(K.RENAMES))
    return A


def _path_of(url):
    m = re.match(r"https?://[^/]+/([^/?#]+)/([^/?#]+)", url)
    return ("%s/%s" % (m.group(1), m.group(2))).lower() if m else None


def issue_sends(A, ym):
    """The issue's sends to the Sluis list (tests and other newsletters left out), one row per send, with its
    Amsterdam send date: a Postwerk campaign sent in batches is one row per batch."""
    y, m = ym
    if ym in [(2026, 4), (2026, 5)]:
        c = A["pw_camp"]
        c = c[(c.send_time.str[3:10] == "%02d-%d" % (m, y)) & (c.list == "Sluis")].copy()
        c["date"] = [dt.datetime.strptime(x, "%d-%m-%Y %H:%M").date() for x in c.send_time]
        return c, "postwerk"
    c = A["rm_camp"]
    c = c[(c.sent_at.str[:7] == "%d-%02d" % (y, m)) & (c.list == "Sluis")].copy()
    c["date"] = [dt.date.fromisoformat(x[:10]) for x in c.sent_at]
    return c, "relaymark"


GATEWAY_LINKS = 4        # a gateway's fetch: one subscriber's clicks on this many links of a send or more ...
GATEWAY_SECONDS = 2      # ... within this many seconds


def readers_clicks(lk, plat, keep_gateways=False):
    """The click rows a reader made: a mail gateway's fetch of every link of a message on delivery is not a click
    through. Relaymark logs each click with its time, so a fetch is a subscriber's clicks on four or more links of
    one send within two seconds; Postwerk reports each subscriber's clicks per link, so a fetch is one click on every
    link of the send, taken off that subscriber's count of each."""
    if keep_gateways or lk.empty:
        return lk
    if plat == "postwerk":
        lk = lk.copy()
        links = lk.groupby("campaign_id").link_label.nunique()
        mine = lk.groupby(["campaign_id", "subscriber_id"]).link_label.nunique()
        fetched = {k for k, v in mine.items() if v == links[k[0]]}
        hit = [(c, s) in fetched for c, s in zip(lk.campaign_id, lk.subscriber_id)]
        lk["clicks"] = lk["clicks"] - np.array(hit, dtype=int)
        return lk[lk.clicks > 0]
    lk = lk.copy()
    lk["t"] = pd.to_datetime(lk.clicked_at, utc=True)
    keep = np.ones(len(lk), dtype=bool)
    pos = {ix: i for i, ix in enumerate(lk.index)}
    for (c, s), grp in lk.groupby(["campaign_id", "subscriber_id"]):
        if grp.link_tag.nunique() < GATEWAY_LINKS:
            continue
        ts = grp.t.values.astype("datetime64[s]").astype(np.int64)
        tags = grp.link_tag.values
        for i, ix in enumerate(grp.index):
            near = np.abs(ts - ts[i]) <= GATEWAY_SECONDS
            if len(set(tags[near])) >= GATEWAY_LINKS:
                keep[pos[ix]] = False
    return lk[keep]


def ask_a(B, A, handling=None):
    """Per issue: the subscribers it was delivered to (the sends of the issue's Tuesday, in batches; the Thursday send
    to the same list goes to subscribers the issue was already delivered to), the subscribers who clicked through to
    the issue's spotlighted repository (distinct subscribers whose own click, not a mail gateway's fetch, is on a link
    whose address is that repository, in any of the issue's sends) and their share of the delivered. handling switches
    the readings the cracker sheet lists."""
    h = handling or {}
    lot = {x["issue"]: x for x in B["lottery"]}
    out = {}
    for ym in ASK_MONTHS:
        issue = K.first_tuesday(*ym)
        rid = lot[issue]["drawn"]
        path = B["repos"][rid]["full_name"].lower()
        sends, plat = issue_sends(A, ym)
        if h.get("sum_all"):
            counted = sends
        elif h.get("first_send"):
            counted = sends.sort_values("date", kind="mergesort").iloc[[0]]
        elif h.get("one_per_campaign"):
            counted = sends.drop_duplicates("campaign_id" if plat == "postwerk" else "campaign_name")
        else:
            counted = sends[sends.date == issue]
        delivered = int(counted.delivered.sum())
        ids = set(counted.campaign_id) if h.get("clicks_counted_only") else set(sends.campaign_id)
        memo = A.setdefault("_readers", {})
        mkey = (tuple(sorted(ids)), plat, bool(h.get("keep_gateways")), GATEWAY_LINKS, GATEWAY_SECONDS)
        if plat == "postwerk":
            lk = A["pw_links"]
            if mkey not in memo:
                memo[mkey] = readers_clicks(lk[lk.campaign_id.isin(ids)], plat, h.get("keep_gateways"))
            lk = memo[mkey]
            if h.get("by_label"):
                hit = lk[lk.link_label == "spotlight"]
            else:
                hit = lk[[_path_of(u) == path for u in lk.url]]
            if h.get("sum_per_link"):
                clickers = int(hit.groupby("link_label").subscriber_id.nunique().sum())
            elif h.get("sum_per_campaign"):
                clickers = int(hit.groupby("campaign_id").subscriber_id.nunique().sum())
            else:
                clickers = int(hit.subscriber_id.nunique())
        else:
            lk = A["rm_links"]
            if mkey not in memo:
                memo[mkey] = readers_clicks(lk[lk.campaign_id.isin(ids)], plat, h.get("keep_gateways"))
            lk = memo[mkey]
            if h.get("by_label"):
                hit = lk[lk.link_tag == "spotlight"]
            elif h.get("feature_only"):
                hit = lk[lk.link_tag == "feature"]
            else:
                hit = lk[[_path_of(u) == path for u in lk.url]]
            if h.get("count_rows"):
                clickers = int(len(hit))
            elif h.get("sum_per_link"):
                clickers = int(hit.groupby("link_tag").subscriber_id.nunique().sum())
            elif h.get("sum_per_campaign"):
                clickers = int(hit.groupby("campaign_id").subscriber_id.nunique().sum())
            else:
                clickers = int(hit.subscriber_id.nunique())
        share = 100.0 * clickers / delivered
        out[ym] = dict(delivered=delivered, clickers=clickers, share=share, share2=K.round_half_up(share, 2),
                       repo=B["repos"][rid]["full_name"])
    return out


# ---------------------------------------------------------------------------------------------------- Ask B

def norm_ref(text):
    t = text.strip()
    t = re.sub(r"^https?://", "", t, flags=re.I)
    t = re.sub(r"^(www\.)?%s/" % re.escape(K.HOST), "", t, flags=re.I)
    t = t.strip("/").lower()
    t = re.sub(r"\.git$", "", t)
    return t


def ask_b(B, A, handling=None):
    """Per shortlisted repository: the people who nominated it (a signed-out response's email hash taken to the
    account whose primary address it is, a merged account to the account it was merged into, a renamed repository's
    old name to its current one) and, for an organisation's repository, the share of them who were not its members
    when they nominated it: a person is a member while any of their accounts is, the accounts later merged into the
    one they nominated from included."""
    h = handling or {}
    prepare(B)
    X = [x for x in B["lottery"] if x["issue"] == OCT_ISSUE][0]
    cands = {B["repos"][rid]["full_name"].lower(): rid for rid in X["shortlist"]}
    bare = {B["repos"][rid]["name"].lower(): rid for rid in X["shortlist"]}
    ren = {} if h.get("no_renames") else {o.lower(): n.lower() for o, n in zip(A["renames"].old_full_name, A["renames"].new_full_name)}
    for o, n in list(ren.items()):
        if n in cands:
            bare.setdefault(o.split("/")[1], cands[n])
    accs = B["accounts"]
    by_hash = {} if h.get("no_hash_map") else dict(zip(accs.primary_email_sha256.dropna(), accs.account_id[accs.primary_email_sha256.notna()]))
    merged = {} if h.get("no_merge_map") else dict(zip(accs.account_id[accs.merged_into.notna()], accs.merged_into.dropna().astype(int)))

    def resolve(aid):
        hops = 0
        while aid in merged and hops < (1 if h.get("single_hop") else 10):
            aid = merged[aid]
            hops += 1
        return aid

    group = {}
    for a in merged:
        group.setdefault(resolve(a), set()).add(a)
    people = {rid: {} for rid in X["shortlist"]}
    for _, r in A["noms"].iterrows():
        ref = norm_ref(r.repository)
        ref = ren.get(ref, ref)
        rid = cands.get(ref) if "/" in ref else bare.get(ref)
        if rid is None:
            continue
        t = K.utc_of(r.submitted_at)
        if r.account_id is not None and not pd.isna(r.account_id):
            key = ("acc", resolve(int(r.account_id)))
        elif r.email_sha256 in by_hash:
            key = ("acc", resolve(int(by_hash[r.email_sha256])))
        else:
            key = ("hash", r.email_sha256)
        people[rid].setdefault(key, []).append(t)
    out = {}
    for rid, ppl in people.items():
        org = owner_of(B, rid)
        n = len(ppl)
        row = dict(name=B["repos"][rid]["name"], people=n, org=org)
        if org:
            inside = 0
            for key, ts in ppl.items():
                if key[0] != "acc":
                    continue
                mine = {key[1]} if h.get("survivor_only") or h.get("members_now") else {key[1]} | group.get(key[1], set())
                if any(is_member(B, org, a, t, current=bool(h.get("members_now"))) for a in mine for t in ts):
                    inside += 1
            row["outside"] = n - inside
            row["share"] = 100.0 * (n - inside) / n
            row["share2"] = K.round_half_up(row["share"], 2)
        out[B["repos"][rid]["name"]] = row
    return out
