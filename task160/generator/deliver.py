"""task160 deliverables: the golden workbook and chart, generator/figures.json and submission.md, every figure computed
from the shipped pack (pack_read.py) and written once."""
from __future__ import annotations

import collections
import datetime as dt
import json
import os
import statistics as st
import sys

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import matplotlib.ticker  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import common as C  # noqa: E402
import pack_read as PR  # noqa: E402
import postprocess as PP  # noqa: E402

D = dt.date
F_BOOK = "winter_pool_block_jan2027.xlsx"
F_CHART = "pool_shifts_winter_2025-26.png"
WEEK0 = D(2025, 11, 3)
N_WEEKS = 21
BLOCK_PREV = (D(2026, 1, 5), D(2026, 2, 1))
REGIONS = ("Southeast", "West and Central", "North")


def r50(x):
    return int(50 * round(x / 50.0))


def compute(target):
    P = PR.Pack(target)
    a27, _ = C.season_bounds(2027)
    b0, b1 = C.BLOCK_2027
    # the closed winters: each member's longest spell, each region's wave centre and the North's lag
    longest = {fid: max((P.spells(fid, *C.season_bounds(w))[0][1] - P.spells(fid, *C.season_bounds(w))[0][0]).days + 1
                        for w in C.CLOSED) for fid in P.fids()}
    lags = [P.region_centre("North", w) - P.region_centre("Southeast", w) for w in C.CLOSED]
    se_move = (max(P.region_centre("Southeast", w) - C.season_bounds(w)[0].toordinal() for w in C.CLOSED)
               - min(P.region_centre("Southeast", w) - C.season_bounds(w)[0].toordinal() for w in C.CLOSED))
    jan = [P.block_draw(P.fids(), w) for w in C.CLOSED]
    north_jan = [P.block_draw(P.fids("North"), w) for w in C.CLOSED]
    carry = [abs(jan[i] - jan[i + 1]) / jan[i + 1] for i in range(len(jan) - 1)]
    # this winter: the south's spells and when each stands down at its own longest closed-winter spell
    south = P.fids("Southeast") + P.fids("West and Central")
    starts, ends = {}, {}
    for fid in south:
        sp = P.spells(fid, a27, C.CENSUS_END)
        assert len(sp) == 1, (fid, sp)
        starts[fid] = sp[0][0]
        ends[fid] = sp[0][0] + dt.timedelta(days=longest[fid] - 1)
    south_block = 0                                   # no southern spell reaches the block weeks
    assert max(ends.values()) < b0
    # the North's wave timed from the Southeast's by every closed winter's lag
    se27 = P.region_centre("Southeast", 2027, upto=C.CENSUS_END)
    north_centres = [se27 + lg for lg in lags]
    half = max(longest[fid] for fid in P.fids("North")) / 2.0
    inside = all(b0.toordinal() <= c - half and c + half <= b1.toordinal() for c in north_centres)
    assert inside
    waves = P.north_waves()
    g = st.mean(waves)
    # the pool's own weeks after each closed winter's Southeast centre, at this winter's offsets
    lo, hi = b0.toordinal() - se27, b1.toordinal() - se27
    post = []
    for w in C.CLOSED:
        c = P.region_centre("Southeast", w)
        a = D.fromordinal(int(round(c + lo)))
        b = D.fromordinal(int(round(c + hi)))
        post.append(sum(P.draw_nights(fid, a, b) for fid in south))
    # the asks
    fA, fB, wk = PR.ask_figures(P, PR.ask_lines(P))
    weeks = []
    for k in range(N_WEEKS):
        c_, i_ = wk[(k, "C")], wk[(k, "I")]
        weeks.append((WEEK0 + dt.timedelta(days=7 * k), c_, i_, c_ + i_))
    busiest = max(weeks, key=lambda x: x[3])
    words = {4: "four", 5: "five", 6: "six", 7: "seven"}
    s0, s1, e1 = min(starts.values()), max(starts.values()), max(ends.values())
    txt = {
        "began": "%d to %d %s" % (s0.day, s1.day, s1.strftime("%B")) if s0.month == s1.month else None,
        "down_by": "%d %s" % (e1.day, e1.strftime("%B")),
        "lag": "%s to %s weeks" % (words[int(min(lags) // 7)], words[int(-(-max(lags) // 7))]),
        "moved": "%s weeks" % words[int(round(se_move / 7.0))],
    }
    assert txt["began"] and txt["lag"] == "five to six weeks" and txt["moved"] == "five weeks"
    return {
        "txt": txt,
        "P": P, "golden": g, "block": r50(g), "waves": waves, "south": south, "starts": starts, "ends": ends,
        "south_block": south_block, "north_block": r50(g), "lags": lags, "se_move": se_move, "jan": jan,
        "north_jan": north_jan, "carry": carry, "se27": se27, "north_centres": north_centres,
        "r0": st.mean(jan), "post": post, "fA": fA, "fB": fB, "weeks": weeks, "busiest": busiest,
    }


def figures(R):
    P = R["P"]
    F = {
        "block_shifts": R["block"],
        "block_weeks": "Monday 4 January to Sunday 31 January 2027",
        "rounding_unit_shifts": 50,
        "clauses_cited": ["3.1", "5.2"],
        "south_expected_draw_block_weeks": R["south_block"],
        "north_expected_draw_block_weeks": R["north_block"],
        "rival_calendar_januaries_averaged": r50(R["r0"]),
        "rival_calendar_januaries_mean": round(R["r0"], 1),
        "rival_calendar_januaries_nearly": int(100 * -(-R["r0"] // 100)),
        "rival_pool_own_post_peak_weeks": r50(st.mean(R["post"])),
        "members": 24, "southern_members": 16, "north_members": 8,
        "closed_winters": 5,
        "carry_min_miss_more_than_a_third": round(min(R["carry"]), 3),
        "southeast_centre_moved_days": R["se_move"],
        "lag_weeks_five_to_six": [5, 6],
        "south_spells_began": R["txt"]["began"] + " 2026",
        "south_all_stand_down_by": R["txt"]["down_by"] + " 2026",
        "busiest_week_total": R["busiest"][3],
        "busiest_week": "%s to %s" % (R["busiest"][0].isoformat(), (R["busiest"][0] + dt.timedelta(days=6)).isoformat()),
        "weeks_charted": N_WEEKS,
        "total_funded": sum(R["fA"][f] for f in R["south"]),
        "total_billed_not_funded": sum(R["fB"][f] for f in R["south"]),
    }
    for fid in R["south"]:
        F["funded_%s" % fid] = R["fA"][fid]
        F["billed_not_funded_%s" % fid] = R["fB"][fid]
    for wk0, c_, i_, t_ in R["weeks"]:
        F["week_%s" % wk0.isoformat()] = {"crestfield": c_, "inland_seas": i_, "total": t_}
    return F


HEAD_FILL = PatternFill("solid", fgColor="DCE6F0")
BLOCK_FILL = PatternFill("solid", fgColor="FBE5C8")
BOLD = Font(name="Calibri", bold=True, size=11)
TITLE = Font(name="Calibri", bold=True, size=13)


def write_book(path, R):
    P = R["P"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Block"
    ws.append(["January 2027 block: %s agency RN shifts, all of it the North's surge wave" % "{:,}".format(R["block"])])
    ws["A1"].font = TITLE
    ws.append(["Block weeks Monday 4 January to Sunday 31 January 2027 (agreement clause 5.1). For the order to "
               "Crestfield Medical Staffing and Inland Seas Nursing on 22 December 2026."])
    ws.append([])
    ws.append(["", "Agency RN shifts"])
    for c in ws[4]:
        c.font = BOLD
        c.fill = HEAD_FILL
    ws.append(["Block (expected draw over the block weeks, nearest 50 shifts, clause 5.2)", R["block"]])
    ws.cell(row=5, column=1).font = BOLD
    ws.cell(row=5, column=2).font = BOLD
    ws.cell(row=5, column=2).fill = BLOCK_FILL
    ws.append(["Expected draw, the 16 Southeast and West and Central members", R["south_block"]])
    ws.append(["Expected draw, the 8 North members (members from 1 November 2026)", R["north_block"]])
    ws.append([])
    tx = R["txt"]
    for line in (
        "Southeast and West and Central: surge plans activated %s 2026; at each member's longest surge "
        "spell in the five closed winters, every one stands down by %s, before the block weeks." % (tx["began"], tx["down_by"]),
        "North: in each closed winter (2021-22 to 2025-26) the North's surge spells came %s after the "
        "Southeast's, while the Southeast's own timing moved by %s. This winter that puts the North's whole "
        % (tx["lag"], tx["moved"]) +
        "wave inside the block weeks, in mid-January.",
        "North draw: each North member's surge nights at its surge plan's agency staffing per level (surge plan "
        "register, 2 November 2026), the eight members' waves averaged over the five closed winters.",
        "Capacity census nights from the capacity system, both layouts, to 20 December 2026; declarations from the "
        "surge declarations log.",
    ):
        ws.append([line])
        ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[ws.max_row].height = 32
    ws.column_dimensions["A"].width = 96
    ws.column_dimensions["B"].width = 18
    for r in (5, 6, 7):
        ws.cell(row=r, column=2).number_format = "#,##0"
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="right")

    ws2 = wb.create_sheet("Members 2025-26")
    ws2.append(["Winter 2025-26 by member, agency RN shifts started 1 November 2025 to 31 March 2026"])
    ws2["A1"].font = TITLE
    ws2.append(["Funded: shifts staffing the member's surge unit that started while its surge plan was active "
                "(clause 3.1). Billed to the pool, not funded: returned to the member's account (clause 6.2)."])
    ws2.append([])
    ws2.append(["Member no.", "Facility id", "Hospital", "Region", "Funded by the pool",
                "Billed to the pool, not funded"])
    for c in ws2[4]:
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    mno = dict(zip(P.roster["facility_id"], P.roster["member_no"]))
    for fid in R["south"]:
        ws2.append([mno[fid], fid, P.name[fid], P.region[fid], R["fA"][fid], R["fB"][fid]])
    ws2.append(["Total", "", "", "", sum(R["fA"][f] for f in R["south"]), sum(R["fB"][f] for f in R["south"])])
    for c in ws2[ws2.max_row]:
        c.font = BOLD
    for row in ws2.iter_rows(min_row=5, max_row=ws2.max_row):
        for c in row[4:]:
            c.number_format = "#,##0"
    for col, w in zip("ABCDEF", [11, 11, 32, 18, 14, 16]):
        ws2.column_dimensions[col].width = w
    ws2.row_dimensions[4].height = 30
    ws2.freeze_panes = "A5"
    ws2.print_title_rows = "4:4"

    ws3 = wb.create_sheet("Weekly 2025-26")
    ws3.append(["Agency RN shifts the pool funded, Monday-to-Sunday weeks by shift start (Eastern Time)"])
    ws3["A1"].font = TITLE
    ws3.append([])
    ws3.append(["Week starting", "Crestfield", "Inland Seas", "Total", "January block week"])
    for c in ws3[3]:
        c.font = BOLD
        c.fill = HEAD_FILL
    for wk0, c_, i_, t_ in R["weeks"]:
        ws3.append([wk0, c_, i_, t_, "yes" if BLOCK_PREV[0] <= wk0 <= BLOCK_PREV[1] else ""])
        ws3.cell(row=ws3.max_row, column=1).number_format = "d mmm yyyy"
    for col, w in zip("ABCDE", [15, 12, 12, 10, 18]):
        ws3.column_dimensions[col].width = w
    ws3.freeze_panes = "A4"
    wb.save(path)
    PP.normalize_ooxml(path, C.PEOPLE["pool_manager"], C.PEOPLE["pool_manager"], "2026-12-21T14:12:00Z",
                       "2026-12-21T21:36:00Z", excel=True, company=C.ORG)


def write_chart(path, R):
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9, "svg.hashsalt": "task160"})
    fig, ax = plt.subplots(figsize=(9.6, 4.8), dpi=130)
    xs = list(range(N_WEEKS))
    cr = [w[1] for w in R["weeks"]]
    isn = [w[2] for w in R["weeks"]]
    tot = [w[3] for w in R["weeks"]]
    k0 = [i for i, w in enumerate(R["weeks"]) if w[0] == BLOCK_PREV[0]][0]
    ax.axvspan(k0 - 0.5, k0 + 3.5, color="#f3e3c3", zorder=0, lw=0, label="January block weeks, 5 Jan to 1 Feb 2026")
    ax.bar(xs, cr, 0.72, label="Crestfield Medical Staffing", color="#2f5d8a", zorder=2)
    ax.bar(xs, isn, 0.72, bottom=cr, label="Inland Seas Nursing", color="#8fb3d6", zorder=2)
    kb = [i for i, w in enumerate(R["weeks"]) if w[0] == R["busiest"][0]][0]
    ax.annotate("{:,}".format(R["busiest"][3]), (kb, R["busiest"][3]), ha="center", va="bottom", fontsize=9,
                fontweight="bold", xytext=(0, 3), textcoords="offset points")
    ax.set_xticks(xs)
    ax.set_xticklabels([w[0].strftime("%d %b").lstrip("0") for w in R["weeks"]], rotation=60, ha="right",
                       fontsize=8)
    ax.set_xlim(-0.6, N_WEEKS - 0.4)
    ax.set_ylim(0, max(tot) * 1.16)
    ax.set_ylabel("Agency RN shifts funded")
    ax.set_xlabel("Week starting (Monday)")
    nz = [w for w in R["weeks"] if w[3]]
    words = {3: "three", 4: "four", 5: "five", 6: "six", 7: "seven", 8: "eight", 9: "nine"}
    ax.set_title("Winter 2025-26: the pool's funded shifts fell in %s weeks from %s, peaking at %s in the week of %s" % (
        words[len(nz)], nz[0][0].strftime("%d %B").lstrip("0"), "{:,}".format(R["busiest"][3]),
        R["busiest"][0].strftime("%d %B").lstrip("0")), fontsize=10.5)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.yaxis.grid(True, color="#e2e2e2", zorder=1)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _p: "{:,.0f}".format(v)))
    hs_, ls_ = ax.get_legend_handles_labels()
    ax.legend(hs_[1:] + hs_[:1], ls_[1:] + ls_[:1], frameon=False, loc="upper right", fontsize=8.5)
    fig.text(0.01, 0.012, "Shifts staffing a member's surge unit that started while its surge plan was active "
             "(agreement clause 3.1), 16 members, weeks of 3 November 2025 to 23 March 2026. TPHA pool office.",
             fontsize=7.3, color="#555555")
    fig.tight_layout(rect=(0, 0.04, 1, 1))
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)


def write_submission(path, R):
    P = R["P"]
    tx = R["txt"]
    L = []
    L.append("## Tags")
    L.append("")
    L.append("**Domain:** Healthcare Delivery & Health Economics (patient flow, beds and workforce optimization: a "
             "hospital association's winter agency nurse pool, each region's respiratory surge timed by its lag behind "
             "the region it spreads from).")
    L.append("**Analytical objective:** Forecasting & Predictive Modeling.")
    L.append("")
    L.append("## 1. Final Recommendation")
    L.append("")
    L.append("**The pool should block-book %s agency RN shifts for the four weeks from Monday 4 January to Sunday 31 "
             "January 2027.** All of it is the eight North members' surge wave, which came %s after the "
             "Southeast's in every closed winter and this winter falls inside the block weeks, while the sixteen "
             "southern members' waves began in early December and end before them. Not the five closed Januaries' "
             "average draw of nearly %s shifts, which reads January by the calendar when the Southeast's surge moved by "
             "%s from winter to winter; not none, the pool's own draw in the same weeks after each closed winter's "
             "peak, which never held a North wave because the North joined the pool on 1 November 2026."
             % ("{:,}".format(R["block"]), tx["lag"], "{:,}".format(int(100 * -(-R["r0"] // 100))), tx["moved"]))
    L.append("")
    L.append("## 2. Critical Components")
    L.append("")
    L.append("1. Expected draw of the sixteen Southeast and West and Central members over the block weeks: **none**.")
    L.append("2. The North's wave falls **wholly inside** the block weeks, in mid-January.")
    L.append("3. Expected draw of the eight North members over the block weeks: **%s** shifts." % "{:,}".format(
        R["north_block"]))
    L.append("")
    L.append("## 3. Step-by-Step Solution")
    L.append("")
    L.append("1. Joined capacity_reports_legacy_2021-07-02_2024-07-01.csv (each row dated by the next morning's report, a "
             "resubmission replacing its original) and capacity_daily_2024-07-01_2026-12-20.csv through "
             "tpha_members_2026-27.xlsx into one series of census nights for each of the 24 members, 1 July 2021 to "
             "20 December 2026.")
    L.append("2. Read each member's surge spells off that series and priced every surge night at its level's agency "
             "staffing in surge_plan_register_2026-27.pdf, the shifts the pool funds under clause 3.1 of "
             "winter_pool_agreement_2026-27.pdf.")
    L.append("3. Back-tested the calendar on the five closed winters, 2021-22 to 2025-26: each winter's draw over its "
             "January block weeks, carried to the next winter, missed that winter's by more than a third every time, "
             "the Southeast's surge moving by %s between winters." % tx["moved"])
    L.append("4. Timed each region's wave in each closed winter by its members' surge spells: the North's came %s "
             "after the Southeast's every winter and never fell in a January block's weeks." % tx["lag"])
    L.append("5. This winter the southern members' surge spells began %s and, at each member's longest "
             "spell in the closed winters, all stand down by %s: the sixteen southern members draw none of "
             "the block weeks." % (tx["began"], tx["down_by"]))
    L.append("6. Carried the closed winters' lag from this winter's Southeast spells: the North's wave falls in "
             "mid-January, wholly inside the four block weeks.")
    L.append("7. Averaged the North's five closed winters' waves at its members' staffing per level: %s shifts to the "
             "nearest 50, the pool's expected draw over the block weeks (clause 5.2), so the block is %s shifts."
             % ("{:,}".format(R["north_block"]), "{:,}".format(R["block"])))
    L.append("")
    L.append("## 4. Deliverable Answers")
    L.append("")
    L.append("### %s" % F_BOOK)
    L.append("")
    L.append("1. Block: %s agency RN shifts, Monday 4 January to Sunday 31 January 2027." % "{:,}".format(R["block"]))
    L.append("2. Agency RN shifts the pool funded between 1 November 2025 and 31 March 2026, and agency shifts worked "
             "there in that time billed to the pool that the pool did not fund, by hospital:")
    for fid in R["south"]:
        L.append("   - %s (%s): %s funded, %s billed to the pool and not funded" % (
            P.name[fid], fid, "{:,}".format(R["fA"][fid]), "{:,}".format(R["fB"][fid])))
    L.append("")
    L.append("### %s" % F_CHART)
    L.append("")
    b = R["busiest"][0]
    L.append("1. Stacked bar chart of the shifts the pool funded in each Monday-to-Sunday week from 3 November 2025 to "
             "29 March 2026, one bar per week split by agency (Crestfield Medical Staffing, Inland Seas Nursing), the "
             "January block weeks of 5 January to 1 February 2026 shaded.")
    L.append("2. Busiest week: %s to %s, %s shifts, labelled on its bar." % (
        b.strftime("%d %B").lstrip("0"), (b + dt.timedelta(days=6)).strftime("%d %B %Y").lstrip("0"),
        "{:,}".format(R["busiest"][3])))
    L.append("")
    open(path, "w", encoding="utf-8").write("\n".join(L))


def main(target, golden_dir, figures_path, submission_path):
    R = compute(target)
    if os.path.isdir(golden_dir):
        for f in os.listdir(golden_dir):
            os.remove(os.path.join(golden_dir, f))
    os.makedirs(golden_dir, exist_ok=True)
    write_book(os.path.join(golden_dir, F_BOOK), R)
    write_chart(os.path.join(golden_dir, F_CHART), R)
    F = figures(R)
    with open(figures_path, "w") as fh:
        json.dump(F, fh, indent=1, sort_keys=True)
        fh.write("\n")
    write_submission(submission_path, R)
    return R


if __name__ == "__main__":
    t = sys.argv[1] if len(sys.argv) > 1 else C.TARGET
    main(t, os.path.join(C.TASK_DIR, "golden"), os.path.join(HERE, "figures.json"),
         os.path.join(C.TASK_DIR, "submission.md"))
