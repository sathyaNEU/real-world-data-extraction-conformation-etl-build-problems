"""task160 independent verifier: reads only the shipped bundle (no generator import) and recomputes the block, its
components, the rivals the write-up names, both member columns and the chart's busiest week on its own code path,
and reads back every format the pack ships with a positive control on the same path. It then reads the goldens and
submission.md and checks they carry the same figures.

    python3 verify.py <target folder>"""
import csv
import datetime as dt
import json
import os
import re
import statistics as st
import sys
from collections import Counter, defaultdict
from zoneinfo import ZoneInfo

import openpyxl
from PIL import Image
from pypdf import PdfReader

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
T = sys.argv[1] if len(sys.argv) > 1 else os.path.join(TASK, "target")
FIG = json.load(open(os.path.join(HERE, "figures.json")))
ET = ZoneInfo("America/New_York")
D = dt.date
fails = 0


def ok(label, cond, detail=""):
    global fails
    print("%s %s %s" % ("PASS" if cond else "FAIL", label, detail if not cond else ""))
    if not cond:
        fails += 1


def path(pattern):
    hits = [n for n in sorted(os.listdir(T)) if re.fullmatch(pattern, n)]
    assert len(hits) == 1, (pattern, hits)
    return os.path.join(T, hits[0])


def pdf_text(p):
    return "\n".join(pg.extract_text() for pg in PdfReader(p).pages)


def r50(x):
    return int(50 * round(x / 50.0))


# ------------------------------------------------------------------------------------------------ roster (xlsx)
wb = openpyxl.load_workbook(path(r"tpha_members_.*\.xlsx"), read_only=True, data_only=True)
rows = [r for r in wb["Members"].iter_rows(values_only=True) if r and r[0]]
ROS = [dict(zip(rows[0], r)) for r in rows[1:]]
wb.close()
ok("xlsx roster: 24 members, member numbers and facility ids unique",
   len(ROS) == 24 and len({r["facility_id"] for r in ROS}) == 24 and len({str(r["member_no"]) for r in ROS}) == 24)
FID_OF = {str(r["member_no"]): r["facility_id"] for r in ROS}
REG = {r["facility_id"]: r["region"] for r in ROS}
UNIT = {r["facility_id"]: r["surge_unit"] for r in ROS}
NAME = {r["facility_id"]: r["hospital"] for r in ROS}
CRS = {r["crestfield_site_code"]: r["facility_id"] for r in ROS}
ISN = {r["inland_seas_facility"]: r["facility_id"] for r in ROS}
ALL = sorted(REG)
NORTH = [f for f in ALL if REG[f] == "North"]
SOUTH = [f for f in ALL if REG[f] != "North"]
SE = [f for f in ALL if REG[f] == "Southeast"]

# ------------------------------------------------------------------------------------------------ the register (pdf)
lines = [x.strip() for x in pdf_text(path(r"surge_plan_register_.*\.pdf")).splitlines()]
GRID, units_ok = {}, True
for r in ROS:
    i = lines.index(r["hospital"])
    units_ok &= lines[i + 1] == r["surge_unit"]
    GRID[r["facility_id"]] = (int(lines[i + 2]), int(lines[i + 3]))
ok("pdf register: every member's surge unit as the roster has it, staffing per level read (Rouge Park 10 and 22)",
   units_ok and GRID["MI-0412"] == (10, 22) and all(b > a for a, b in GRID.values()))

# ------------------------------------------------------------------------------------------------ the agreement (pdf)
AG = " ".join(pdf_text(path(r"winter_pool_agreement_.*\.pdf")).split())
ok("pdf agreement: clauses 1.2, 1.4, 3.1, 5.1, 5.2 and 6.2 read back",
   "A Member is a hospital listed in Schedule A, from its Pool start date." in AG
   and "A shift is one nurse for one scheduled shift, 07:00 to 19:00 or 19:00 to 07:00." in AG
   and "The Pool funds the agency RN shifts that staff a Member's surge unit and start while the Member's surge plan is active." in AG
   and "The January Block covers every Member's surge unit for the four weeks from the first Monday in January: "
       "Monday 4 January to Sunday 31 January 2027." in AG and "rounded to the nearest 50 shifts" in AG
   and "returned to the Member's account" in AG)
ok("pdf agreement: Schedule A lists all 24 roster hospitals", all(r["hospital"] in AG for r in ROS))
RATES = " ".join(pdf_text(path(r"agency_rate_schedule_.*\.pdf")).split())
ok("pdf rate schedule read back (Crestfield block rate 88.50)", "Block rate, per hour" in RATES and "88.50" in RATES)

# ------------------------------------------------------------------------------------------------ capacity census (csv)
CEN = defaultdict(dict)
leg = {}
with open(path(r"capacity_reports_legacy_.*\.csv"), newline="") as fh:
    for x in csv.DictReader(fh):
        k = (x["member_no"], x["report_date"])
        if k in leg and leg[k]["submission"] == "R":
            continue
        leg[k] = x
for (mno, rd), x in leg.items():
    night = D.fromisoformat(rd) - dt.timedelta(days=1)
    CEN[FID_OF[mno]][night] = {"": 0, "S1": 1, "S2": 2}[x["surge_status"].strip()]
n_leg = sum(len(v) for v in CEN.values())
with open(path(r"capacity_daily_.*\.csv"), newline="") as fh:
    for x in csv.DictReader(fh):
        d = D.fromisoformat(x["census_date"])
        assert d not in CEN[x["facility_id"]]
        CEN[x["facility_id"]][d] = int(x["surge_level"])
span = (D(2026, 12, 20) - D(2021, 7, 1)).days + 1
ok("csv capacity: both layouts stitch, every member every census night 1 July 2021 to 20 December 2026",
   all(len(CEN[f]) == span and min(CEN[f]) == D(2021, 7, 1) and max(CEN[f]) == D(2026, 12, 20) for f in ALL)
   and n_leg == 24 * ((D(2024, 6, 30) - D(2021, 7, 1)).days + 1))

# ------------------------------------------------------------------------------------------------ declarations (csv)
IV = defaultdict(list)
with open(path(r"surge_declarations_.*\.csv"), newline="") as fh:
    dec = sorted(csv.DictReader(fh), key=lambda x: (x["facility_id"], x["declared_at"]))
cur = {}
for x in dec:
    t = dt.datetime.strptime(x["declared_at"], "%Y-%m-%d %H:%M")
    f = x["facility_id"]
    if f in cur:
        IV[f].append((cur[f][0], t, cur[f][1]))
        del cur[f]
    if int(x["level_to"]) > 0:
        cur[f] = (t, int(x["level_to"]))
for f, (t, lv) in cur.items():
    IV[f].append((t, dt.datetime(2100, 1, 1), lv))
ok("csv declarations: times only at 06:45 and 14:45", all(x["declared_at"][-5:] in ("06:45", "14:45") for x in dec))


def level(f, t):
    for a, b, lv in IV[f]:
        if a <= t < b:
            return lv
    return 0


# ------------------------------------------------------------------------------------------------ the call
CLOSED = [2022, 2023, 2024, 2025, 2026]


def season(w):
    return D(w - 1, 11, 1), D(w, 3, 31)


def block_weeks(w):
    d = D(w, 1, 1)
    while d.weekday() != 0:
        d += dt.timedelta(days=1)
    return d, d + dt.timedelta(days=27)


def draw(f, a, b):
    return sum(GRID[f][lv - 1] for d, lv in CEN[f].items() if a <= d <= b and lv)


def draw_by_shift(f, a, b):
    n, d = 0, a
    while d <= b:
        for hh in (7, 19):
            lv = level(f, dt.datetime.combine(d, dt.time(hh)))
            n += GRID[f][lv - 1] // 2 if lv else 0
        d += dt.timedelta(days=1)
    return n


def spells(f, a, b):
    out, s = [], None
    d = a
    while d <= b:
        lv = CEN[f].get(d, 0)
        if lv and s is None:
            s = d
        if not lv and s is not None:
            out.append((s, d - dt.timedelta(days=1)))
            s = None
        d += dt.timedelta(days=1)
    if s is not None:
        out.append((s, b))
    return out


def centre(fids, w, upto=None):
    a, b = season(w)
    b = min(b, upto) if upto else b
    mids = []
    for f in fids:
        sp = spells(f, a, b)
        assert len(sp) == 1, (f, w, sp)
        mids.append(sp[0][0].toordinal() + (sp[0][1] - sp[0][0]).days / 2.0)
    return st.median(mids)


WC = [f for f in ALL if REG[f] == "West and Central"]
PULL = D(2026, 12, 20)
waves = [sum(draw(f, *season(w)) for f in NORTH) for w in CLOSED]
g = st.mean(waves)
block = r50(g)
ok("call: the North's mean per-wave draw at the register's staffing rounds to the block of 600",
   block == 600 == FIG["block_shifts"] == FIG["north_expected_draw_block_weeks"], (waves, g))
ok("bin: the golden at least a quarter bin inside its cell (edge distance from the figure)", 25 - abs(g - block) >= 12.5,
   25 - abs(g - block))
wv_s = [sum(draw_by_shift(f, *season(w)) for f in NORTH) for w in CLOSED]
mid3 = sorted(waves)[1:4]
ok("call: shift by shift at the declared times, the median wave, the last three and the middle three file 600",
   r50(st.mean(wv_s)) == 600 and r50(st.median(waves)) == 600 and r50(st.mean(waves[-3:])) == 600
   and r50(st.mean(mid3)) == 600, (wv_s, waves))
jan = [sum(draw(f, *block_weeks(w)) for f in ALL) for w in CLOSED]
ok("rival: the five closed Januaries' draws average about 1,850",
   r50(st.mean(jan)) == FIG["rival_calendar_januaries_averaged"] == 1850, jan)
ok("rival: shift by shift, the closed Januaries still average about 1,850",
   r50(st.mean(sum(draw_by_shift(f, *block_weeks(w)) for f in ALL) for w in CLOSED)) == 1850)
carry = [abs(jan[i] - jan[i + 1]) / jan[i + 1] for i in range(4)]
carry_f = [abs(jan[i] - jan[i + 1]) / jan[i] for i in range(4)]
ok("back-test: each January carried to the next winter missed it by a quarter or more, either way it is scaled",
   min(carry) >= 0.25 and min(carry_f) >= 0.25, (carry, carry_f))
ok("the North's surge nights in no closed January block", all(sum(draw(f, *block_weeks(w)) for f in NORTH) == 0
                                                              for w in CLOSED))
lag_wc = [centre(NORTH, w) - centre(WC, w) for w in CLOSED]
lag_se = [centre(NORTH, w) - centre(SE, w) for w in CLOSED]
ok("timing: the North's waves 30 to 34 days after West and Central's in every closed winter",
   all(30 <= x <= 34 for x in lag_wc), lag_wc)
ok("timing: the Southeast's lead over the North moved by four weeks", round((max(lag_se) - min(lag_se)) / 7.0) == 4,
   lag_se)
off = [centre(SE, w) - season(w)[0].toordinal() for w in CLOSED]
ok("timing: the Southeast's own timing moved by about five weeks",
   round((max(off) - min(off)) / 7.0) == 5 and max(off) - min(off) == FIG["southeast_centre_moved_days"], off)
loo_wc = [abs(st.mean(lag_wc[:i] + lag_wc[i + 1:]) - lag_wc[i]) for i in range(5)]
loo_se = [abs(st.mean(lag_se[:i] + lag_se[i + 1:]) - lag_se[i]) for i in range(5)]
ok("back-test: the lag behind West and Central from the other four winters within two days of each winter's own; "
   "the lag behind the Southeast misses a winter by more than three weeks", max(loo_wc) < 2 and max(loo_se) > 21,
   (loo_wc, loo_se))
a27 = D(2026, 11, 1)
st27 = {f: spells(f, a27, PULL) for f in SOUTH}
s_on = [v[0][0] for v in st27.values()]
wc_on = [st27[f][0][0] for f in WC]
ok("this winter: one surge spell at each southern member, begun 2 to 9 December; West and Central's 3 to 9 December",
   all(len(v) == 1 for v in st27.values()) and min(s_on) == D(2026, 12, 2) and max(s_on) == D(2026, 12, 9)
   and min(wc_on) == D(2026, 12, 3) and max(wc_on) == D(2026, 12, 9)
   and FIG["south_spells_began"] == "2 to 9 December 2026" and FIG["west_and_central_spells_began"] == "3 to 9 December 2026")
longest = {f: max((spells(f, *season(w))[0][1] - spells(f, *season(w))[0][0]).days + 1 for w in CLOSED) for f in ALL}
ends = {f: st27[f][0][0] + dt.timedelta(days=longest[f] - 1) for f in SOUTH}
ok("this winter: at each member's longest closed-winter spell every southern member stands down by 29 December",
   max(ends.values()) == D(2026, 12, 29) and FIG["south_all_stand_down_by"] == "29 December 2026")
b0, b1 = D(2027, 1, 4), D(2027, 1, 31)
half = max(longest[f] for f in NORTH) / 2.0


def inside(c):
    return b0.toordinal() <= c - half and c + half <= b1.toordinal()


wc27 = centre(WC, 2027, upto=PULL)
on_wc = [st.median(spells(f, *season(w))[0][0].toordinal() for f in WC) for w in CLOSED]
on27 = st.median(d.toordinal() for d in wc_on)
lag_on = [centre(NORTH, w) - on_wc[i] for i, w in enumerate(CLOSED)]
ok("this winter: at every closed winter's lag behind West and Central, from its centre at the pull or its spells' "
   "onset, the North's wave falls wholly inside the block weeks, in mid-January",
   all(inside(wc27 + x) for x in lag_wc) and all(inside(on27 + x) for x in lag_on)
   and D(2027, 1, 10) <= D.fromordinal(int(round(on27 + st.mean(lag_on)))) <= D(2027, 1, 21),
   (wc27, lag_wc, on27, lag_on))
ok("this winter: no North member in surge at the pull", all(not spells(f, a27, PULL) for f in NORTH))


def shifted(fids, c27, cw, w):
    k = int(round(c27 - cw))
    return sum(draw(f, b0 - dt.timedelta(days=k), b1 - dt.timedelta(days=k)) for f in fids)


al = st.mean(shifted(ALL, wc27, centre(WC, w), w) for w in CLOSED)
pl = st.mean(shifted(NORTH, wc27 + st.mean(lag_wc), centre(NORTH, w), w) for w in CLOSED)
ok("call: the whole census aligned on West and Central's centre, and the North placed at its lag, file 600",
   r50(al) == 600 and r50(pl) == 600, (al, pl))
se27 = centre(SE, 2027, upto=PULL)
se_t = [st.mean(shifted(NORTH, se27 + agg(lag_se), centre(NORTH, w), w) for w in CLOSED) for agg in (st.mean, st.median)]
ok("rival: timed from the Southeast, the North's wave leaves the block a small part of the 600",
   max(se_t) < 0.25 * g and max(FIG["rival_north_timed_from_the_southeast"]) < 0.25 * g, se_t)
ok("component: the sixteen southern members' expected draw over the block weeks is none",
   FIG["south_expected_draw_block_weeks"] == 0 and max(ends.values()) < b0)

# ------------------------------------------------------------------------------------------------ the asks
LINES = []
with open(path(r"crestfield_member_shifts_.*\.csv"), newline="") as fh:
    cr = list(csv.DictReader(fh))
by_ts = defaultdict(list)
for x in cr:
    by_ts[x["timesheet_no"]].append(x)
hours_ok = True
for ts, xs in by_ts.items():
    xs.sort(key=lambda z: z["shift_start_utc"])
    hours_ok &= abs(sum(float(z["hours"]) for z in xs) - 12.0) < 1e-9
    t = dt.datetime.fromisoformat(xs[0]["shift_start_utc"].replace("Z", "+00:00")).astimezone(ET).replace(tzinfo=None)
    LINES.append((CRS[xs[0]["client_site"]], "C", t, xs[0]["unit"], xs[0]["bill_to"] == "TPHA POOL"))
ok("csv Crestfield: every timesheet bills 12 hours, nights on two lines split at local midnight",
   hours_ok and any(len(v) == 2 for v in by_ts.values()))
night_ok, n_isn = True, 0
for fn in (r"inland_seas_member_shifts_2025-26\.xlsx", r"inland_seas_member_shifts_2026-27\.xlsx"):
    wb = openpyxl.load_workbook(path(fn), read_only=True, data_only=True)
    for ws in wb.worksheets:
        vals = [r for r in ws.iter_rows(values_only=True)]
        hi_ = [i for i, r in enumerate(vals) if r and r[0] == "Facility"][0]
        head = vals[hi_]
        for r in vals[hi_ + 1:]:
            if not r or not r[0]:
                continue
            x = dict(zip(head, r))
            if ws.title == "Timesheet detail":
                cin = dt.datetime.strptime(x["Clock in"], "%m/%d/%Y %H:%M")
                if x["Shift"] == "N":
                    night_ok &= cin.date() == x["Payroll date"].date() - dt.timedelta(days=1)
                continue
            pdate = x["Payroll date"].date()
            t = (dt.datetime.combine(pdate - dt.timedelta(days=1), dt.time(19)) if x["Shift"] == "N"
                 else dt.datetime.combine(pdate, dt.time(7)))
            LINES.append((ISN[x["Facility"]], "I", t, x["Unit"], x["Bill to"] == "TPHA Pool"))
            n_isn += 1
    wb.close()
ok("xlsx Inland Seas: every queried night clocked in the evening before its payroll date", night_ok)
A, B, WK, BW = Counter(), Counter(), Counter(), Counter()
W0 = D(2025, 11, 3)
for f, ag, t, u, pool in LINES:
    if not (D(2025, 11, 1) <= t.date() <= D(2026, 3, 31)):
        continue
    k = (t.date() - W0).days // 7
    fund = u == UNIT[f] and level(f, t) > 0
    if fund:
        A[f] += 1
        if 0 <= k < 21:
            WK[(k, ag)] += 1
    elif pool:
        B[f] += 1
    if pool and 0 <= k < 21:
        BW[k] += 1
ok("ask A: the shifts the pool funded at each of the 16 members", all(A[f] == FIG["funded_%s" % f] for f in SOUTH),
   [(f, A[f]) for f in SOUTH])
ok("ask B: the shifts billed to the pool it did not fund at each of the 16 members",
   all(B[f] == FIG["billed_not_funded_%s" % f] for f in SOUTH), [(f, B[f]) for f in SOUTH])
ok("asks: no figure nil and no North member in last winter's billing", all(A[f] > 0 and B[f] > 0 for f in SOUTH)
   and not any(A[f] or B[f] for f in NORTH))
ok("asks: the table's totals", sum(A[f] for f in SOUTH) == FIG["total_funded"]
   and sum(B[f] for f in SOUTH) == FIG["total_billed_not_funded"])
tot = Counter()
for (k, ag), v in WK.items():
    tot[k] += v
kb = max(range(21), key=lambda k: (tot[k], -k))
ok("ask C: the busiest week by shifts started, its funded total and the agency shifts starting that week billed to "
   "the pool", tot[kb] == FIG["busiest_week_total"] and BW[kb] == FIG["busiest_week_billed_to_pool"]
   and (W0 + dt.timedelta(days=7 * kb)).isoformat() == FIG["busiest_week"][:10]
   and (W0 + dt.timedelta(days=7 * kb + 6)).isoformat() == FIG["busiest_week"][-10:], (kb, tot[kb], BW[kb]))
ok("ask C: the busiest week clear of the next busiest", tot[kb] - max(tot[k] for k in range(21) if k != kb) >= 20,
   sorted(tot.values())[-3:])

# ------------------------------------------------------------------------------------------------ the other files
ww = json.load(open(path(r"wastewater_results_.*\.json")))


def ww_mean(reg, a, b):
    v = [r["concentration_copies_per_l"] / r["pmmov_copies_per_l"] for r in ww["results"]
         if r["sewershed_id"].startswith("WW-" + reg) and r["target"] == "Influenza A" and a <= r["sample_date"] <= b
         and r["concentration_copies_per_l"] is not None]
    return st.mean(v)


ok("json wastewater: 19,000 results or more; the North's influenza flat at the pull while the Southeast's has risen",
   len(ww["results"]) >= 19000 and ww_mean("NO", "2026-12-04", "2026-12-17") / ww_mean("NO", "2026-10-01", "2026-11-15") < 1.3
   and ww_mean("SE", "2026-12-04", "2026-12-17") / ww_mean("SE", "2026-10-01", "2026-11-15") > 10)
wb = openpyxl.load_workbook(path(r"resp_admissions_outlook_.*\.xlsx"), read_only=True, data_only=True)
otx = " ".join(str(c) for ws in wb.worksheets for r in ws.iter_rows(values_only=True) for c in r if c is not None)
wb.close()
ok("xlsx outlook read back (issued 18 December 2026)", "18 December 2026" in otx)
wb = openpyxl.load_workbook(path(r"rn_workforce_survey_.*\.xlsx"), read_only=True, data_only=True)
n_sv = sum(1 for ws in wb.worksheets for r in ws.iter_rows(values_only=True) if r and r[0])
wb.close()
ok("xlsx workforce survey read back (distractor)", n_sv > 100, n_sv)
wb = openpyxl.load_workbook(path(r"january_block_orders_.*\.xlsx"), read_only=True, data_only=True)
bo = [r for ws in wb.worksheets for r in ws.iter_rows(values_only=True) if r and r[0]]
wb.close()
bo_h = [dict(zip(bo[0], r)) for r in bo[1:]]
ok("xlsx block orders: five filed Januaries, each after the first the last January's draw to the nearest 50, split "
   "between the agencies, the 2026-27 row not yet filed",
   len(bo_h) == 6 and all(bo_h[i]["block_shifts"] == r50(jan[i - 1]) for i in range(1, 5))
   and all(x["crestfield_shifts"] + x["inland_seas_shifts"] == x["block_shifts"] for x in bo_h[:5])
   and bo_h[5]["block_shifts"] is None, [x["block_shifts"] for x in bo_h])
wb = openpyxl.load_workbook(path(r"data_dictionary\.xlsx"), read_only=True, data_only=True)
dd = [r for ws in wb.worksheets for r in ws.iter_rows(values_only=True) if r and r[0]]
wb.close()
ok("xlsx dictionary read back: Inland Seas pays every shift under the day it ends, in both winters' files",
   sum(1 for r in dd if r[1] == "Payroll date" and "the day the shift ends" in str(r[3])) == 2 and len(dd) > 20)
with open(path(r"extract_manifest_.*\.csv"), newline="") as fh:
    man = list(csv.DictReader(fh))
ok("csv manifest names every other shipped file", {m["file"] for m in man} == set(os.listdir(T)) - {
    os.path.basename(path(r"extract_manifest_.*\.csv"))})
files = os.listdir(T)
big = max(sum(1 for _ in open(os.path.join(T, f))) - 1 for f in files if f.endswith(".csv"))
ok("input gates: 10 or more files, 3 or more formats, a file of 25,000 rows or more",
   len(files) >= 10 and len({f.rsplit(".", 1)[1] for f in files}) >= 3 and big >= 25000, (len(files), big))

# ------------------------------------------------------------------------------------------------ goldens and write-up
GD = os.path.join(TASK, "golden")
wb = openpyxl.load_workbook(os.path.join(GD, "winter_pool_block_jan2027.xlsx"), read_only=True, data_only=True)
bl = [r for r in wb["Block"].iter_rows(values_only=True)]
mem = [r for r in wb["Members 2025-26"].iter_rows(values_only=True) if r and r[0]]
names = wb.sheetnames
wb.close()
ok("golden workbook: the block and its two components", bl[4][1] == 600 and bl[5][1] == 0 and bl[6][1] == 600
   and names == ["Block", "Members 2025-26"], names)
mrow = {r[1]: r for r in mem if r[1] in A or r[1] in SOUTH}
ok("golden workbook: the member table equals the recomputed columns",
   all(mrow[f][4] == A[f] and mrow[f][5] == B[f] for f in SOUTH) and len(mrow) == 16)
im = Image.open(os.path.join(GD, "pool_shifts_winter_2025-26.png"))
ok("golden chart: a rendered PNG", im.format == "PNG" and im.size[0] > 1000 and im.size[1] > 500, im.size)
sub = open(os.path.join(TASK, "submission.md"), encoding="utf-8").read()
ok("submission: the call, both rivals and every member line carry the recomputed figures",
   "block-book 600 agency RN shifts" in sub and "average draw of about 1,850 shifts" in sub
   and "timed from the Southeast, whose lead over the North moved by four weeks" in sub
   and "began %s" % FIG["west_and_central_spells_began"][:-5] in sub
   and all("(%s): %s funded, %s billed to the pool and not funded" % (f, "{:,}".format(A[f]), B[f]) in sub
           for f in SOUTH) and "{:,} shifts funded".format(tot[kb]) in sub
   and "billed to the pool: {:,}".format(BW[kb]) in sub)
ok("submission: no figure in blocks 1 to 3 finer than 50 shifts beside the block (about 1,850 and about 600)",
   not re.search(r"\b(?:602|60[0-9]\.\d|1,86\d)\b", sub.split("## 4.")[0]))
ok("submission: no em or en dash", chr(0x2014) not in sub and chr(0x2013) not in sub)
cited = set(re.findall(r"\b[\w.-]+\.(?:csv|xlsx|pdf|json|png)\b", sub))
ok("submission: every file it names is shipped or a golden", bool(cited) and all(
    c in os.listdir(T) or c in os.listdir(GD) for c in cited), sorted(cited))

print("verify: %d failed" % fails)
sys.exit(1 if fails else 0)
