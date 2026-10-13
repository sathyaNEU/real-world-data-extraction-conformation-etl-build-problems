"""Writers for the data files under target/."""
import csv
import datetime as dt

import openpyxl
from openpyxl.styles import Font

import sim as S
import world as W

FILES = W.FILES


def _csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\r\n")
        w.writerow(header)
        n = 0
        for r in rows:
            w.writerow(r)
            n += 1
    return n


def _money(x):
    return "%.2f" % x


def write_claims(claims, path):
    rows = []
    order = sorted(claims, key=lambda c: (c["paid"], c["fill"], c["claim_id"]))
    newid = {c["claim_id"]: 2600000000 + 1000 + k * 3 + (c["claim_id"] % 3) for k, c in enumerate(order)}
    for c in order:
        q = c["qty"]
        qty = ("%.1f" % q) if c["product"].liquid else "%d" % q
        days = 30 if c["product"].liquid else abs(int(q))
        rows.append(("%010d" % newid[c["claim_id"]], c["txn"],
                     "%010d" % newid[c["original"]] if c["original"] else "",
                     c["fill"].isoformat(), c["paid"].isoformat(), c["member"], c["pharmacy"], c["ndc"]["ndc11"], qty,
                     days, _money(c["ingredient"]), _money(c["fee"]), {"UC": "U&C"}.get(c["basis"], c["basis"])))
    return _csv(path, ("claim_id", "transaction_code", "original_claim_id", "fill_date", "paid_date", "member_id",
                       "pharmacy_npi", "ndc", "quantity_dispensed", "days_supply", "ingredient_cost_paid",
                       "dispensing_fee_paid", "basis_of_cost"), rows)


def ndc_description(n):
    p = n["product"]
    return ("%s %s %s" % (p.name, p.strength, p.form)).upper()


def write_aac(rates, path):
    last_month = S.WMONTHS[-1]
    rows = []
    for r in sorted(rates, key=lambda r: (r["ndc11"], r["effective"])):
        open_ended = r["end"] >= S.month_end(last_month)
        rows.append((r["ndc11"], ndc_description(r["ndc"]), "%.5f" % r["rate"], r["unit"], r["effective"].isoformat(),
                     "" if open_ended else r["end"].isoformat()))
    return _csv(path, ("ndc", "ndc_description", "aac_per_unit", "pricing_unit", "effective_date", "end_date"), rows)


def _book(sheets):
    wb = openpyxl.Workbook()
    first = True
    for title, header, rows, widths in sheets:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = title
        ws.append(list(header))
        for c in range(1, len(header) + 1):
            ws.cell(row=1, column=c).font = Font(bold=True)
        for r in rows:
            ws.append(list(r))
        for i, wdt in enumerate(widths):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = wdt
        ws.freeze_panes = "A2"
    return wb


def marketing_start(n):
    if n.get("launch"):
        return n["launch"] - dt.timedelta(days=int(S.rng("ms", n["ndc11"]).integers(3, 20)))
    g = S.rng("ms", n["ndc11"])
    return dt.date(int(g.integers(2008, 2024)), int(g.integers(1, 13)), int(g.integers(1, 28)))


def write_drugs(prods, path):
    rows = []
    for p in prods:
        for n in p.ndcs:
            lab = S.LAB[n["labeler"]][0]
            rows.append((n["ndc10"], ("%s %s %s" % (p.name, p.strength, p.form)).upper(), p.name, p.strength, p.form,
                         "ORAL", lab, n["desc"].replace("%s" % p.form.upper(), p.form.upper()), marketing_start(n),
                         p.gpc))
    rows.sort(key=lambda r: (r[1], r[0]))
    wb = _book([("Drug file", ("NDC", "Product description", "Generic name", "Strength", "Dosage form", "Route",
                               "Labeler", "Package description", "Marketing start", "GPC"), rows,
                 (15, 46, 26, 14, 26, 8, 36, 58, 15, 8))])
    ws = wb["Drug file"]
    for row in ws.iter_rows(min_row=2, min_col=9, max_col=9):
        for c in row:
            c.number_format = "yyyy-mm-dd"
    wb.save(path)
    return len(rows)


CLASS_CODES = {}


def class_code(name):
    if name not in CLASS_CODES:
        CLASS_CODES[name] = "C%03d" % (100 + 7 * (len(CLASS_CODES) + 1))
    return CLASS_CODES[name]


def write_classes(prods, path):
    names = W.SHORTLIST_NAMES + sorted({p.cls for p in prods if not p.shortlisted})
    for nm in names:
        class_code(nm)
    crow = [(class_code(nm), nm, "Listed" if nm in W.SMAC_LISTED else "Not listed") for nm in names]
    prow = sorted(((p.gpc, p.name, p.strength, p.form, class_code(p.cls)) for p in prods), key=lambda r: (r[4], r[1], r[2]))
    wb = _book([("Classes", ("Class code", "Class", "SMAC status"), crow, (12, 40, 14)),
                ("Products", ("GPC", "Generic name", "Strength", "Dosage form", "Class code"), prow, (8, 30, 14, 28, 12))])
    wb["Classes"].insert_rows(1)
    wb["Classes"]["A1"] = "Department of Health Coverage drug class list, in force from 1 January 2027"
    wb["Classes"]["A1"].font = Font(bold=True)
    wb["Classes"].freeze_panes = "A3"
    wb.save(path)
    return len(prow)


def write_summary(prior, path):
    agg = {}
    for r in prior:
        k = (r["month"], r["ndc"]["ndc11"])
        a = agg.setdefault(k, dict(r, claims=0, units=0.0, ingredient=0.0))
        a["claims"] += r["claims"]
        a["units"] += r["units"]
        a["ingredient"] += r["ingredient"]
    rows = []
    for k in sorted(agg):
        r = agg[k]
        rows.append((r["month"].strftime("%Y-%m"), r["ndc"]["ndc11"], ndc_description(r["ndc"]), r["label"], r["claims"],
                     round(r["units"], 1) if r["product"].liquid else int(round(r["units"])), round(r["ingredient"], 2)))
    tot = {}
    for r in prior:
        k = r["month"].strftime("%Y-%m")
        t = tot.setdefault(k, [0, 0.0, 0.0])
        t[0] += r["claims"]
        t[1] += r["units"]
        t[2] += r["ingredient"]
    trows = [(k, v[0], round(v[1], 1), round(v[2], 2)) for k, v in sorted(tot.items())]
    wb = _book([("Utilization", ("Fill month", "NDC", "Drug", "Drug class", "Claims", "Units", "Ingredient cost paid"),
                 rows, (11, 13, 46, 34, 9, 12, 18)),
                ("Control totals", ("Fill month", "Claims", "Units", "Ingredient cost paid"), trows, (11, 10, 14, 18))])
    wb.save(path)
    return len(rows)


DICTIONARY = [
    (FILES["claims"], "(file)", "One row per claim transaction on the Program's fee-for-service pharmacy claims for the "
                                "candidate classes, for fills from 1 April 2026 to 31 March 2027, paid through 20 April 2027."),
    (FILES["claims"], "claim_id", "The transaction's identifier."),
    (FILES["claims"], "transaction_code", "B1 for a paid claim; B2 for the reversal of a paid claim."),
    (FILES["claims"], "original_claim_id", "On a reversal, the claim_id of the claim it reverses."),
    (FILES["claims"], "fill_date", "The date the prescription was filled."),
    (FILES["claims"], "paid_date", "The date the transaction was adjudicated for payment."),
    (FILES["claims"], "member_id", "The Program's member identifier."),
    (FILES["claims"], "pharmacy_npi", "The dispensing pharmacy's National Provider Identifier."),
    (FILES["claims"], "ndc", "The eleven-digit NDC dispensed."),
    (FILES["claims"], "quantity_dispensed", "Metric decimal quantity: each for solid forms, millilitres for liquids."),
    (FILES["claims"], "days_supply", "Days supply submitted by the pharmacy."),
    (FILES["claims"], "ingredient_cost_paid", "Ingredient cost the Program paid, in dollars."),
    (FILES["claims"], "dispensing_fee_paid", "Professional dispensing fee the Program paid, in dollars."),
    (FILES["claims"], "basis_of_cost", "The basis the ingredient was paid on: AAC (the State AAC rate in effect on the "
                                       "fill date), WAC, or U&C (the pharmacy's usual and customary charge)."),
    (FILES["aac"], "(file)", "One row per State AAC rate in effect at any time from 1 April 2026 to 31 March 2027."),
    (FILES["aac"], "ndc", "The eleven-digit NDC the rate prices."),
    (FILES["aac"], "aac_per_unit", "The rate in dollars per pricing unit."),
    (FILES["aac"], "pricing_unit", "EA (each) or ML (millilitre)."),
    (FILES["aac"], "effective_date", "First day the rate is in effect."),
    (FILES["aac"], "end_date", "Last day the rate is in effect; blank while in effect."),
    (FILES["summary"], "Utilization", "FY2026 fee-for-service generic utilization by fill month and NDC, net of reversals, "
                                      "prepared 24 April 2026."),
]


def write_dictionary(path):
    return _csv(path, ("file", "field", "description"), DICTIONARY)


def write_transmittal(path, counts):
    lines = [
        "Corvane Pharmacy Solutions: data extract for the State MAC update proposal",
        "Prepared 21 April 2027 for the Department of Health Coverage, Pharmacy Program",
        "",
        "Files:",
        "  %s  (%s rows; fills 1 April 2026 to 31 March 2027, transactions paid through 20 April 2027)" % (
            FILES["claims"], "{:,}".format(counts["claims"])),
        "  %s  (%s rows; the State AAC rates in effect over the same months, as published to 1 April 2027)" % (
            FILES["aac"], "{:,}".format(counts["aac"])),
        "  %s  (%s NDCs; the Department's drug file as of 15 April 2027)" % (FILES["drugs"], "{:,}".format(counts["drugs"])),
        "  %s  (the Department's drug class list)" % FILES["classes"],
        "  %s  (%s rows; FY2026 utilization summary, as delivered 24 April 2026)" % (
            FILES["summary"], "{:,}".format(counts["summary"])),
        "  %s, %s" % (FILES["mco"], FILES["pdl"]),
        "",
        "Claims are limited to the twelve candidate classes the Department named on 8 March 2027.",
        "Contact: Todd Allen, account director.",
    ]
    with open(path, "w", encoding="utf-8", newline="") as fh:
        fh.write("\r\n".join(lines) + "\r\n")
    return len(lines)


def write_mco(prods, path):
    """the managed care plans' pharmacy encounters by class and quarter (the SMAC list applies to fee-for-service
    claims only)"""
    rows = []
    plans = ("Harborview Health Plan", "Summit Community Care", "Prairie Rose Health")
    for cls in W.SHORTLIST_NAMES:
        for q in ("2026Q2", "2026Q3", "2026Q4", "2027Q1"):
            for pl in plans:
                g = S.rng("mco", cls, q, pl)
                enc = int(g.integers(2500, 21000))
                units = int(enc * float(g.uniform(32, 58)))
                paid = round(units * float(g.uniform(0.05, 0.4)), 2)
                rows.append((pl, cls, q, enc, units, "%.2f" % paid))
    return _csv(path, ("plan", "drug_class", "quarter", "encounters", "units", "plan_paid_amount"), rows)


def write_pdl(prods, path):
    rows = []
    for p in prods:
        if not p.shortlisted:
            continue
        g = S.rng("pdl", *p.key)
        rows.append((p.cls, ("%s %s %s" % (p.name, p.strength, p.form)), "Preferred" if g.random() < 0.8 else "Non-preferred",
                     dt.date(2027, 1, 1)))
    wb = _book([("PDL", ("Class", "Drug", "PDL status", "Effective"), rows, (34, 52, 14, 12))])
    for row in wb["PDL"].iter_rows(min_row=2, min_col=4, max_col=4):
        for c in row:
            c.number_format = "yyyy-mm-dd"
    wb.save(path)
    return len(rows)
