"""The golden deliverables, computed from the shipped files only: smac_class_addition_2027-07.docx (the administrator's
proposal to the Department: the class, each candidate's ingredient spread, a chart, how the spread was measured) and
smac_candidate_classes_fy2026.xlsx (each candidate class's claims, units and ingredient cost paid over the twelve
months of fills before the window, classes as the current list defines them), with figures.json."""
import io
import json
import math
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import calc as C        # noqa: E402
import pack as PK       # noqa: E402
import world as W       # noqa: E402

DOCX, XLSX = W.DELIVERABLES


def k1000(v):
    return int(round(v / 1000.0)) * 1000


def compute(target):
    cl, dr, cls = C.frame(target)
    X = {"routes": {}}
    for r in C.ROUTES:
        X["routes"][r] = C.spreads(target, r, cl, dr, cls)[0]
    g = X["routes"]["golden"]
    X["golden"] = g
    X["rank"] = C.ranking(g)
    X["default_rank"] = C.ranking(X["routes"]["default"])
    X["call"] = X["rank"][0]
    X["lo"] = {c: min(X["routes"][r][c] for r in C.GOLDEN) for c in W.SHORTLIST_NAMES}
    X["hi"] = {c: max(X["routes"][r][c] for r in C.GOLDEN) for c in W.SHORTLIST_NAMES}
    net = cl.groupby("cls").agg(claims=("transaction_code", lambda s: int((s == "B1").sum() - (s == "B2").sum())),
                                ingredient=("ingredient_cost_paid", "sum"))
    X["fy27"] = {c: (int(net.loc[c, "claims"]), float(net.loc[c, "ingredient"])) for c in W.SHORTLIST_NAMES}
    X["ask"] = C.ask(target)
    X["ask_label"] = C.ask(target, by_label=True, late=False)
    X["ask_summary"] = C.ask(target, late=False)
    X["ask_overlap"] = C.ask(target, overlap=True)
    rt = C.rates(target)
    by_ndc, by_grp = C.in_effect(rt, dr)
    X["premium"] = C.premium(by_grp, cls, pooled=True)[None]
    unpriced = cl[[by_ndc.get((n, m)) is None for n, m in zip(cl.ndc, cl.month)]]
    X["unpriced_share"] = float(unpriced.ingredient_cost_paid.sum() / cl.ingredient_cost_paid.sum())
    X["unpriced_carton_share"] = float(unpriced[unpriced.pkg == "U"].ingredient_cost_paid.sum() /
                                       unpriced.ingredient_cost_paid.sum())
    allx = C.claims(target)
    early = allx[allx.fill < W.WINDOW[0]]
    X["extract_rows"], X["window_rows"] = len(allx), len(cl)
    X["fy26_later"] = int((early.paid > W.SUMMARY_PAID_THROUGH).sum())
    X["fy26_overlap"] = int((early.paid <= W.SUMMARY_PAID_THROUGH).sum())
    return X


def money(v):
    return "${:,.0f}".format(v)


def chart_png(X):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams["font.family"] = "DejaVu Sans"
    order = list(reversed(X["rank"]))
    vals = [X["golden"][c] / 1000.0 for c in order]
    cols = ["#1f4e5f" if c == X["call"] else "#a9bfc8" for c in order]
    fig, ax = plt.subplots(figsize=(6.4, 3.9), dpi=150)
    ax.barh(range(len(order)), vals, color=cols, height=0.62)
    for k, v in enumerate(vals):
        ax.text(v + 0.5, k, "$%dk" % int(round(v)), va="center", ha="left", fontsize=7.5)
    ax.set_yticks(range(len(order)))
    ax.set_yticklabels(order, fontsize=8)
    ax.set_xlabel("Ingredient spread, fills April 2026 to March 2027 (thousand dollars)", fontsize=8)
    ax.set_xlim(0, max(vals) * 1.15)
    ax.tick_params(axis="x", labelsize=7.5)
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    ax.set_title("Ingredient spread by candidate class", fontsize=9.5, loc="left")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def write_docx(X, path):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    call, second = X["rank"][0], X["rank"][1]
    h = d.add_paragraph()
    r = h.add_run("Corvane Pharmacy Solutions")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x5F)
    d.add_paragraph("Anticonvulsants proposed for the 1 July 2027 State MAC update")
    for k, v in (("To", "%s, Pharmacy Director, %s" % (W.PEOPLE["director"], W.DEPARTMENT)),
                 ("From", "%s, MAC Pricing Lead" % W.PEOPLE["pricing"]),
                 ("Copy", "%s, Account Director" % W.PEOPLE["account"]),
                 ("Date", "30 April 2027"),
                 ("Re", "Proposal under Attachment C, section C.4")):
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        a = p.add_run(k + ": ")
        a.bold = True
        p.add_run(v)
    d.add_paragraph()
    p = d.add_paragraph()
    p.add_run("We propose %s (%s) for the State MAC list from 1 July 2027. " % (call, PK.class_code(call))).bold = True
    p.add_run("Its ingredient spread over fills from April 2026 to March 2027 was %s, the largest of the twelve candidate "
              "classes; %s were next at %s." % (money(k1000(X["golden"][call])), second.lower() if second != call else second,
                                                money(k1000(X["golden"][second]))))
    t = d.add_table(rows=1, cols=4)
    t.style = "Light Grid Accent 1"
    hdr = ("Candidate class", "Claims", "Ingredient cost paid", "Ingredient spread")
    for i, x in enumerate(hdr):
        t.rows[0].cells[i].text = x
    for c in X["rank"]:
        row = t.add_row().cells
        row[0].text = c
        row[1].text = "{:,}".format(X["fy27"][c][0])
        row[2].text = money(X["fy27"][c][1])
        row[3].text = money(k1000(X["golden"][c]))
        for j in (1, 2, 3):
            row[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    src = d.add_paragraph("Source: Corvane claims extract of 21 April 2027 (claims paid through 20 April 2027); State AAC "
                          "rates as published to 1 April 2027. Claims net of reversals. Spreads rounded to the nearest "
                          "thousand dollars; the classes are compared on unrounded figures (C.4).")
    src.runs[0].font.size = Pt(8.5)
    d.add_paragraph()
    d.add_picture(io.BytesIO(chart_png(X)), width=Inches(6.2))
    cap = d.add_paragraph("Figure 1. Ingredient spread by candidate class, fills April 2026 to March 2027, net of "
                          "reversals; the proposed class in dark.")
    cap.runs[0].italic = True
    d.add_paragraph().add_run("How the spread was measured").bold = True
    up = X["unpriced_share"] * 100
    d.add_paragraph(
        "Claims are the Program's fee-for-service claims for fills from 1 April 2026 to 31 March 2027, reversals netted, "
        "each NDC placed in its class through the drug file and the class list in force from 1 January 2027. A claim's "
        "units are costed at the State AAC rate in effect for its NDC on the fill date. %.0f per cent of the candidate "
        "classes' ingredient cost was paid on NDCs with no rate in effect, %.0f per cent of it on unit-dose cartons and "
        "the rest mostly on new entrants' bottles; those units are costed at the median rate in effect that month for "
        "the priced NDCs of the same product and package. Where a product had no priced carton, its cartons are costed "
        "at its bottles' rate times the carton-to-bottle ratio of the products with both priced (about %.2f)." % (
            up, X["unpriced_carton_share"] * 100, X["premium"]))
    d.add_paragraph(
        "A carton costs more per unit than the bottle of the same product. Costed at the product's bottle rate, the "
        "cartons would carry a spread they do not have, and %s would appear first at %s." % (
            X["default_rank"][0].lower(), money(k1000(X["routes"]["default"][X["default_rank"][0]]))))
    d.add_paragraph("The candidate classes' claims, units and ingredient cost paid for fills from April 2025 to March 2026 "
                    "are in the attached workbook for the budget: the FY2026 utilization summary, which holds the "
                    "payments made through 20 April 2026, with the later payments on those fills from this extract, "
                    "classes as the current list defines them.")
    d.add_paragraph()
    d.add_paragraph("%s\nMAC Pricing Lead" % W.PEOPLE["pricing"])
    d.core_properties.author = W.PEOPLE["pricing"]
    d.core_properties.title = "State MAC update, 1 July 2027: proposed class"
    d.save(path)


def write_xlsx(X, path):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Candidate classes"
    ws.append(["Candidate classes: fills April 2025 to March 2026, claims paid through 20 April 2027 (classes as the "
               "list in force from 1 January 2027)"])
    ws["A1"].font = Font(bold=True)
    ws.append(["Class code", "Class", "Claims", "Units dispensed", "Ingredient cost paid"])
    for c in range(1, 6):
        ws.cell(row=2, column=c).font = Font(bold=True)
    for c in W.SHORTLIST_NAMES:
        cl, un, ing = X["ask"][c]
        ws.append([PK.class_code(c), c, cl, int(round(un)), round(ing)])
    tot = [sum(X["ask"][c][i] for c in W.SHORTLIST_NAMES) for i in range(3)]
    ws.append([None, "All candidate classes", tot[0], int(round(tot[1])), round(tot[2])])
    last = ws.max_row
    for c in range(2, 6):
        ws.cell(row=last, column=c).font = Font(bold=True)
    for col, wdt in zip("ABCDE", (11, 30, 10, 16, 20)):
        ws.column_dimensions[col].width = wdt
    for row in ws.iter_rows(min_row=3, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = "#,##0"
    ws.freeze_panes = "A3"
    nt = wb.create_sheet("Notes")
    for line in ("Notes",
                 "Claims are fee-for-service claims for fills from 1 April 2025 to 31 March 2026, net of reversals.",
                 "Source: the FY2026 utilization summary of 24 April 2026 (transactions paid through 20 April 2026), "
                 "with the later payments on those fills from the claims extract of 21 April 2027.",
                 "Classes as the Department's drug class list in force from 1 January 2027 defines them, each NDC placed "
                 "through the drug file of April 2027.",
                 "Units are metric decimal quantities: each for solid forms, millilitres for liquids. Ingredient cost "
                 "paid to the nearest dollar."):
        nt.append([line])
    nt["A1"].font = Font(bold=True)
    nt.column_dimensions["A"].width = 110
    wb.save(path)


def figures(X):
    F = {"call": X["call"], "default_call": X["default_rank"][0], "runner_up": X["rank"][1],
         "golden_margin": round(X["golden"][X["rank"][0]] / X["golden"][X["rank"][1]], 2),
         "default_margin": round(X["routes"]["default"][X["default_rank"][0]] /
                                 X["routes"]["default"][X["default_rank"][1]], 2),
         "carton_premium": round(X["premium"], 2), "unpriced_share_pct": round(X["unpriced_share"] * 100),
         "unpriced_carton_share_pct": round(X["unpriced_carton_share"] * 100), "classes": {},
         "extract_rows": "{:,}".format(X["extract_rows"]), "window_rows": "{:,}".format(X["window_rows"]),
         "fy26_later_payments": "{:,}".format(X["fy26_later"]), "fy26_overlap": "{:,}".format(X["fy26_overlap"]),
         "carton_premium_2dp": "%.2f" % X["premium"], "classes_ranked": X["rank"], "update": "1 July 2027",
         "default_spread_call": "{:,}".format(k1000(X["routes"]["default"][X["default_rank"][0]]))}
    for c in W.SHORTLIST_NAMES:
        cl, un, ing = X["ask"][c]
        lcl, lun, ling = X["ask_label"][c]
        F["classes"][c] = {
            "code": PK.class_code(c), "spread": k1000(X["golden"][c]), "spread_exact": round(X["golden"][c]),
            "spread_range": [int(math.floor(X["lo"][c] / 1000.0)) * 1000, int(math.ceil(X["hi"][c] / 1000.0)) * 1000],
            "default_spread": k1000(X["routes"]["default"][c]), "claims_fy27": X["fy27"][c][0],
            "ingredient_fy27": round(X["fy27"][c][1]),
            "claims_fy26": cl, "units_fy26": int(round(un)), "ingredient_fy26": int(round(ing)),
            "claims_fy26_label": lcl, "units_fy26_label": int(round(lun)), "ingredient_fy26_label": int(round(ling))}
    return F


def main(target, out, figpath):
    os.makedirs(out, exist_ok=True)
    X = compute(target)
    import postprocess as PP
    write_docx(X, os.path.join(out, DOCX))
    write_xlsx(X, os.path.join(out, XLSX))
    PP.normalize_ooxml(os.path.join(out, DOCX), W.PEOPLE["pricing"], W.PEOPLE["pricing"], "2027-04-28T14:05:00Z",
                       "2027-04-30T20:12:41Z", title="State MAC update, 1 July 2027: proposed class")
    PP.normalize_ooxml(os.path.join(out, XLSX), W.PEOPLE["pricing"], W.PEOPLE["pricing"], "2027-04-29T13:40:00Z",
                       "2027-04-30T19:58:06Z", excel=True)
    F = figures(X)
    with open(figpath, "w") as fh:
        json.dump(F, fh, indent=2)
        fh.write("\n")
    return X


if __name__ == "__main__":
    T = os.path.dirname(HERE)
    main(os.path.join(T, "target"), os.path.join(T, "golden"), os.path.join(HERE, "figures.json"))
