"""Cuts Vennholt Energy's files from the world: the battery register, the two lease operator apps (the route app to 30 June,
the check-in app from 1 July, with two routes on both from 15 June), the release event register, the SCADA alarm journal, daily
production, the tank vapour analyses, the maintenance work orders, the leak surveys and their tracker, and the files the
return does not draw on."""
import csv
import datetime as dt
import json

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

import world as W

R = None
SCHED = {"twice": "Twice daily (weekdays)", "daily": "Daily", "mwf": "Mon Wed Fri"}


def rng():
    global R
    if R is None:
        R = np.random.default_rng(W.SEED + 1)
    return R


def code_at(code, day):
    """the battery code a system wrote on a day: the re-licensed battery kept its old code until the change"""
    old, new, when = W.RENAMED
    if code == new and day < when:
        return old
    return code


def lstr(h, sec=False):
    return W.local_of(h).strftime("%Y-%m-%d %H:%M:%S" if sec else "%Y-%m-%d %H:%M")


def ustr(h, fmt="%Y-%m-%dT%H:%M:%SZ"):
    return W.utc_of(h).strftime(fmt)


def ostr(h):
    """the check-in app's export: local time with its UTC offset"""
    return W.utc_of(h).astimezone(W.TZ).replace(microsecond=0).isoformat()


def write_csv(path, rows, fields, crlf=False):
    with open(path, "w", newline="", encoding="utf-8") as f:
        wr = csv.DictWriter(f, fieldnames=fields, lineterminator="\r\n" if crlf else "\n")
        wr.writeheader()
        for r in rows:
            wr.writerow(r)
    return len(rows)


def write_xlsx(path, sheets, widths=None, freeze="A2"):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, header, rows in sheets:
        ws = wb.create_sheet(name)
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows:
            ws.append(r)
        ws.freeze_panes = freeze
        for i, h in enumerate(header, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(9, min(46, len(str(h)) + 3)))
    wb.save(path)


# ---------------------------------------------------------------- the battery register
def renamed_new(w):
    old, new, when = W.RENAMED
    return new


def apply_rename(w):
    """the world names the battery by its current code; the re-licensing gave BT-0907 a new code on 17 August"""
    old, new, when = W.RENAMED
    if old in w.by:
        b = w.by.pop(old)
        b["code"] = new
        w.by[new] = b
        for d in (w.visits, w.prod, w.analyses):
            d[new] = d.pop(old)
        for e in w.events:
            if e["bat"] == old:
                e["bat"] = new
        for x in w.alarms:
            if x["bat"] == old:
                x["bat"] = new
        w.outages = [(new if c == old else c, a, z) for c, a, z in w.outages]
        for d in w.devices:
            if d["bat"] == old:
                d["bat"] = new
        for x in w.wells:
            if x["bat"] == old:
                x["bat"] = new
        w.fuel_ch4[new] = w.fuel_ch4.pop(old)
        w.casing_ch4[new] = w.casing_ch4.pop(old)
        for lk in w.ldar["leaks"]:
            if lk["bat"] == old:
                lk["bat"] = new
        for k in ("surveys", "turnaround"):
            if old in w.ldar[k]:
                w.ldar[k][new] = w.ldar[k].pop(old)
        to = lambda c: new if c == old else c
        for k in ("contract", "closure"):
            w.ldar[k] = {(to(c), d, kind): v for (c, d, kind), v in w.ldar[k].items()}
        w.ldar["sessions"] = {(to(c), d): [(to(a), x, y) for a, x, y in ks] for (c, d), ks in w.ldar["sessions"].items()}


def register(path, w):
    rows = []
    for b in sorted(w.bats, key=lambda b: b["code"]):
        rows.append(dict(battery=b["code"], facility_licence=b["lic"], route=b["run27"], visit_schedule=SCHED[b["cls"]],
                         saturday_visits="Y" if b["sat"] and b["cls"] != "mwf" else "N", tanks=b["tanks"],
                         tank_pressure_telemetry="Y" if b["tel"] else "N", vapour_combustor="Y" if b["comb"] else "N",
                         former_battery=W.RENAMED[0] if b["code"] == W.RENAMED[1] else "", status="Active"))
    return write_csv(path, rows, list(rows[0]))


# ---------------------------------------------------------------- the two lease operator apps
def visit_rows(w):
    """every visit with what each app saw of it: the route app's taps to the minute, the check-in app's geofence to the second"""
    r = rng()
    ev_at = {(e["bat"], round(e["arr"], 6)): e for e in w.events}
    out = []
    for b in w.bats:
        for v in w.visits[b["code"]]:
            e = ev_at.get((b["code"], round(v["arr"], 6)))
            out.append(dict(b=b, v=v, e=e, day=W.day_of(v["arr"]),
                            fl_in=v["arr"] + r.uniform(0, 2) / 60, fl_out=v["dep"] - r.uniform(0, 2) / 60,
                            sl_in=v["arr"] - r.uniform(.3, 1.5) / 60, sl_out=v["dep"] + r.uniform(.3, 1.5) / 60,
                            manual=r.random() < .06, split=r.random() < .04, gap=r.uniform(1, 6) / 60,
                            cut=r.uniform(.25, .75), remark=r.random()))
    out.sort(key=lambda x: (x["v"]["arr"], x["b"]["code"]))
    return out


def comb_word(x, sl):
    if not x["b"]["comb"]:
        return None if sl else ""
    if x["e"] is not None and x["e"]["kind"] == "CP":
        return "relit" if sl else "Out - relit"
    return "lit" if sl else "Lit"


REMARKS_E = {"TH": ["hatch open T%d, closed", "T%d hatch was open - shut and latched", "found T%d thief hatch up"],
             "CP": ["combustor pilot out, relit", "ECD out - relit pilot", "pilot was out"],
             "PV": ["PVRV T%d not seated, reseated", "vent valve stuck open T%d", "T%d relief valve open - fixed"]}
REMARKS = [("hauler on site", 31), ("tank level high, called truck", 19), ("pump off on arrival - restarted", 14),
           ("chem pump topped up", 11), ("met mechanic", 7), ("heater checked", 6), ("road rough", 5),
           ("flowline drained", 3), ("gate left open, closed it", 2), ("fence down on lease", 1), ("cattle on lease", 1)]


def remark(u):
    """a free remark from the operator's short list, the common ones far more often"""
    tot = sum(k for _, k in REMARKS)
    acc = 0
    for t, k in REMARKS:
        acc += k
        if u * tot < acc:
            return t
    return REMARKS[-1][0]


def fieldlog(path, w, vr):
    r = rng()
    rows = []
    n = 418204
    for x in vr:
        if not (W.FIRST <= x["day"] <= dt.date(2026, 6, 30)):
            continue
        b, v, e = x["b"], x["v"], x["e"]
        rem = ""
        if e is not None and x["remark"] < .45:
            t = REMARKS_E[e["kind"]][int(x["remark"] * 100) % 3]
            rem = t % e["tank"] if "%d" in t else t
        elif x["remark"] > .93:
            rem = remark((x["remark"] - .93) / .07)
        fl_out = max(x["fl_out"], x["fl_in"] + 3 / 60)
        rows.append(dict(VisitNo="FL%06d" % n, Site=code_at(b["code"], x["day"]), Route=b["route"],
                         LeaseOperator=v["user"], Arrived=lstr(x["fl_in"]), Departed=lstr(fl_out),
                         TanksGauged=b["tanks"] if v["kind"] == "sched" or r.random() < .7 else max(1, b["tanks"] - 1),
                         Combustor=comb_word(x, False), Remarks=rem))
        n += 1 + (1 if r.random() < .004 else 0)
    return write_csv(path, rows, list(rows[0]), crlf=True)


def site_ids(w):
    r = np.random.default_rng(W.SEED + 7)
    ids = r.choice(np.arange(10000, 99999), size=len(w.bats), replace=False)
    return {b["code"]: "S-%05d" % int(i) for b, i in zip(sorted(w.bats, key=lambda b: b["code"]), ids)}


def sitelink(path, w, vr, sites):
    r = rng()
    n = 1000317
    recs = []
    for x in vr:
        b, v = x["b"], x["v"]
        start = W.PILOT if b["route"] in W.PILOT_ROUTES else W.MIGRATE
        if x["day"] < start:
            continue
        base = dict(site=sites[b["code"]], run=W.run_on(b, x["day"]), user=v["user"],
                    tanks_gauged=b["tanks"] if v["kind"] == "sched" or r.random() < .7 else max(1, b["tanks"] - 1),
                    combustor=comb_word(x, True))
        if x["manual"]:
            a = round((x["fl_in"]) * 60) / 60
            c = max(a + 4 / 60, round(x["fl_out"] * 60) / 60)
            recs.append(dict(base, check_in=ostr(a), check_out=ostr(c), method="manual"))
        elif x["split"] and x["sl_in"] + x["cut"] * (x["sl_out"] - x["sl_in"]) + x["gap"] < x["sl_out"] - 1 / 60:
            mid = x["sl_in"] + x["cut"] * (x["sl_out"] - x["sl_in"])
            recs.append(dict(base, check_in=ostr(x["sl_in"]), check_out=ostr(mid), method="geofence"))
            recs.append(dict(base, check_in=ostr(mid + x["gap"]), check_out=ostr(x["sl_out"]), method="geofence",
                             tanks_gauged=0, combustor=None if not b["comb"] else "lit"))
        else:
            recs.append(dict(base, check_in=ostr(x["sl_in"]), check_out=ostr(x["sl_out"]), method="geofence"))
    recs.sort(key=lambda d: (dt.datetime.fromisoformat(d["check_in"]), d["site"]))
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        for d in recs:
            d = dict(id="SL%07d" % n, **d)
            n += 1
            f.write(json.dumps({k: d[k] for k in ("id", "site", "run", "user", "check_in", "check_out", "method",
                                                  "tanks_gauged", "combustor")}) + "\n")
    return len(recs)


def sites_csv(path, w, sites):
    r = np.random.default_rng(W.SEED + 8)
    rows = []
    for code in sorted(sites, key=lambda c: sites[c]):
        b = w.by[code]
        rows.append(dict(site_id=sites[code], site_name="%s battery" % code, battery=code,
                         geofence_radius_m=int(r.choice([100, 120, 150, 150, 200])),
                         created="2026-%02d-%02d" % (5 + int(r.integers(0, 2)), 1 + int(r.integers(0, 28)))))
    return write_csv(path, rows, list(rows[0]))


# ---------------------------------------------------------------- the release event register
TEXT = {"TH": ["Thief hatch open T%d", "Hatch open on tank %d", "T%d hatch found open", "Open thief hatch, tank %d",
               "Hatch lid open T%d", "Tank %d thief hatch open"],
        "CP": ["Combustor pilot out", "ECD pilot not lit", "Combustor flame out", "Pilot out on vapour combustor",
               "VCU pilot out"],
        "PV": ["PVRV stuck open T%d", "Relief valve not reseated, tank %d", "PRV open and venting T%d",
               "Tank %d vent valve open", "PVRV pallet hung open T%d"]}
ACTION = {"TH": ["Closed and latched", "Closed hatch, checked gasket", "Shut and latched lid", "Closed"],
          "CP": ["Relit pilot", "Relit, checked igniter", "Drained knockout, relit", "Pilot relit"],
          "PV": ["Reseated PVRV", "Freed pallet, reseated", "Replaced PVRV seal", "Reseated valve, checked set"]}


# days from the find to the office entry: most within the week, some at the month-end catch-up
ENTRY_LAG = [0] * 9 + [1] * 11 + [2] * 6 + [3] * 4 + [4, 4, 5, 5, 6, 7, 8, 9, 11, 13, 16, 19, 23, 27]


def event_register(path, w):
    r = rng()
    rows = []
    evs = [e for e in w.events if W.day_of(e["found"]).year == W.YEAR]
    order = sorted(evs, key=lambda e: (e["found"] + r.uniform(0, 30), e["bat"]))
    users = {}
    for b in w.bats:
        for v in w.visits[b["code"]]:
            users[(b["code"], round(v["arr"], 6))] = v["user"]
    for i, e in enumerate(order, 1):
        d = W.day_of(e["found"])
        t = TEXT[e["kind"]][int(r.integers(0, len(TEXT[e["kind"]])))]
        txt = t % e["tank"] if "%d" in t else t
        ent = min(d + dt.timedelta(days=int(r.choice(ENTRY_LAG))), dt.date(2027, 1, 7))
        if d < W.MIGRATE:
            fs, cs = lstr(e["found"]), lstr(e["closed"])
        else:
            fs, cs = ostr(round(e["found"] * 60) / 60), ostr(round(e["closed"] * 60) / 60)
        rows.append(["RE26-%04d" % i, code_at(e["bat"], d), e["tank"], fs, txt, ACTION[e["kind"]][int(r.integers(0, 4))],
                     cs, users[(e["bat"], round(e["arr"], 6))], ent])
    hdr = ["Event No", "Battery", "Tank", "Found", "Event", "Action", "Closed", "Reported by", "Entered"]
    write_xlsx(path, [("Register 2026", hdr, rows)], widths={"Event": 34, "Action": 30, "Found": 17, "Closed": 17,
                                                             "Entered": 11, "Event No": 11})
    wb = openpyxl.load_workbook(path)
    ws = wb["Register 2026"]
    for row in ws.iter_rows(min_row=2):
        row[8].number_format = "yyyy-mm-dd"
    ws.auto_filter.ref = "A1:I%d" % (len(rows) + 1)
    wb.save(path)
    return len(rows)


# ---------------------------------------------------------------- the SCADA alarm journal
def scada_name(code):
    """SCADA points kept the battery's old code through its re-licensing"""
    return W.RENAMED[0] if code == W.RENAMED[1] else code


def latency(u):
    """seconds from the RTU's alarm to the host's journal over the radio: most reports by exception within seconds, a
    share waiting for the next poll, a few after retries"""
    if u < .72:
        return .4 + u / .72 * 5.6
    if u < .96:
        return 6 + (u - .72) / .24 * 34
    return 40 + (u - .96) / .04 * 200


def shift_op(t, cover):
    """the control room works twelve-hour shifts, 06:00 to 18:00 and back, two crews on days and two on nights,
    each on four days and off four; a crew short on a shift is covered by the crew it hands over to"""
    loc = W.local_of(t)
    day = loc.hour >= 6 and loc.hour < 18
    k = (loc - dt.timedelta(hours=6)).date().toordinal() // 4 % 2
    pair = ("ctrl01", "ctrl02") if day else ("ctrl03", "ctrl04")
    op = pair[k]
    if any(a <= loc.date() <= b for a, b in LEAVE.get(op, ())):
        op = pair[1 - k]
    elif cover == 0 and loc.date().toordinal() % 9 == 0:
        op = ("ctrl03", "ctrl04")[k] if day else ("ctrl01", "ctrl02")[k]
    return op


LEAVE = {"ctrl01": [(dt.date(2026, 7, 6), dt.date(2026, 7, 26))], "ctrl02": [(dt.date(2026, 2, 9), dt.date(2026, 2, 22))],
         "ctrl03": [(dt.date(2026, 8, 10), dt.date(2026, 8, 30)), (dt.date(2026, 12, 21), dt.date(2027, 1, 3))],
         "ctrl04": [(dt.date(2026, 3, 16), dt.date(2026, 3, 29))]}


def alarm_journal(path, w):
    r = rng()
    rows = []
    ops = ["ctrl01", "ctrl02", "ctrl03", "ctrl04"]
    for x in w.alarms:
        pt = "KES/%s/TK/PIT-101" % scada_name(x["bat"])
        if x["state"] == "ACK":
            t = max(x["t"], max(y["rx"] for y in w.alarms_by_lift[(x["bat"], x["lift"])])) + r.uniform(5, 300) / 3600
            cover = int(r.integers(0, 4))
            rows.append((t, pt, "ACK", "", "", "Acknowledged", shift_op(t, cover)))
            continue
        rx = x["rx"] if x["rx"] > x["t"] else x["t"] + latency(r.uniform(0, 1)) / 3600
        rtu = W.utc_of(x["t"]).strftime("%Y-%m-%d %H:%M:%S")
        if x["state"] == "ALM":
            rows.append((rx, pt, "ALARM", "%.1f" % x["val"], "kPa", "Tank pressure high | RTU %s" % rtu, ""))
        else:
            rows.append((rx, pt, "NORMAL", "%.1f" % x["val"], "kPa", "Tank pressure normal | RTU %s" % rtu, ""))
    for code, a, c in w.outages:
        pt = "KES/%s/RTU" % scada_name(code)
        rows.append((a + r.uniform(60, 240) / 3600, pt, "COMM FAIL", "", "", "RTU communication failure", ""))
        rows.append((c, pt, "COMM OK", "", "", "RTU communication restored", ""))
    lo, hi = W.h_of(dt.datetime(2025, 12, 15)), W.h_of(dt.datetime(2027, 1, 13))
    rows = sorted((x for x in rows if lo <= x[0] < hi), key=lambda x: (x[0], x[1], x[2]))
    out = [dict(Time=W.utc_of(t).strftime("%Y-%m-%d %H:%M:%S"), Point=p, State=s, Value=v, Units=u, Message=m,
                Operator=o) for t, p, s, v, u, m, o in rows]
    return write_csv(path, out, ["Time", "Point", "State", "Value", "Units", "Message", "Operator"])


# ---------------------------------------------------------------- production and analyses
def production(path, w):
    rows = []
    by_day = {}
    for b in w.bats:
        for p in w.prod[b["code"]]:
            by_day.setdefault(p["day"], []).append((code_at(b["code"], p["day"]), p))
    for day in sorted(by_day):
        for code, p in sorted(by_day[day]):
            rows.append(dict(Battery=code, ProdDate=day.isoformat(), OilVol_m3="%.1f" % p["oil"],
                             WaterVol_m3="%.1f" % p["water"], GasVol_e3m3="%.2f" % p["gas"],
                             HoursOn="%.1f" % p["hours"]))
    return write_csv(path, rows, list(rows[0]))


def vapour_analyses(path, w):
    r = rng()
    rows = []
    n = 25100
    for b in sorted(w.bats, key=lambda b: b["code"]):
        for a in w.analyses[b["code"]]:
            rest = 100 - a["ch4"]
            parts = np.array([r.uniform(1, 3), r.uniform(.5, 2), r.uniform(10, 16), r.uniform(8, 14), r.uniform(1.5, 3),
                              r.uniform(3, 6), r.uniform(1, 2), r.uniform(1, 2), r.uniform(.8, 2.4)])
            parts = parts / parts.sum() * rest
            p = [round(float(v), 2) for v in parts]
            p[2] = round(rest - sum(p) + p[2], 2)
            lab = "Quellmark Laboratories" if a["date"].year == 2025 or r.random() < .6 else "Tavistrom Analytical"
            rows.append([code_at(b["code"], a["date"]), a["date"], "QL-%05d" % n if lab[0] == "Q" else "TA%06d" % (n * 3),
                         "Tank vapour (flash)", a["gor"], p[0], p[1], a["ch4"]] + p[2:] + [lab])
            n += int(r.integers(1, 9))
    rows.sort(key=lambda x: (x[1], x[0]))
    hdr = ["Battery", "Sampled", "Lab ref", "Sample point", "Flash GOR (m3/m3)", "N2", "CO2", "C1", "C2", "C3", "iC4",
           "nC4", "iC5", "nC5", "C6+", "Laboratory"]
    write_xlsx(path, [("Results", hdr, rows)], widths={"Sample point": 18, "Flash GOR (m3/m3)": 11, "Laboratory": 18})
    wb = openpyxl.load_workbook(path)
    ws = wb["Results"]
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "yyyy-mm-dd"
    wb.save(path)
    return len(rows)


# ---------------------------------------------------------------- maintenance
PM = ["Monthly separator inspection", "Quarterly pumpjack PM", "Treater firetube inspection", "Tank thief hatch gasket check",
      "Combustor igniter service", "PVRV inspection", "Flowline pig run", "Wellhead greasing", "Fire extinguisher check",
      "Gas detector bump test"]
CM = ["Replaced dump valve trim", "Repaired flowline leak", "Replaced packing on wellhead stuffing box",
      "Pump motor replaced", "Heater tube cleaned", "Replaced tank level gauge float", "Repaired fence",
      "Replaced pressure gauge", "Freed stuck dump valve", "Replaced combustor flame arrestor"]


def received(w, x):
    """the day a tag reached the tracker: in 2026 the second crew's tablet syncs at the end of its week"""
    return synced(x["survey"]) if w.ldar_crew[(x["bat"], x["survey"])] == "ogi08" else x["survey"]


def work_orders(path, w):
    """the maintenance system's work orders completed in 2026, numbered as they were raised. The field's own
    maintainers complete theirs on the day of the work; a contract work order is completed when the foreman approves
    its field tickets. The reference is the notification a work order was raised on, its maintenance plan, or the
    contractor's field ticket"""
    import asks as A
    r = rng()
    rq = np.random.default_rng(W.SEED + 45)
    rows = []          # battery, type, description, day of the work, day raised, contract, the leak program's key
    for b in w.bats:
        for m in range(1, 13):
            d = dt.date(2026, m, 1) + dt.timedelta(days=int(r.integers(0, 27)))
            if r.random() < .05:
                continue
            rows.append([b["code"], "Preventive", PM[0], d, dt.date(2026, m, 1), False, None])
            if m % 3 == 1:
                rows.append([b["code"], "Preventive", PM[int(r.integers(1, len(PM)))],
                             d + dt.timedelta(days=int(r.integers(0, 3))), dt.date(2026, m, 1), False, None])
        for _ in range(int(r.poisson(3.2))):
            d = dt.date(2026, 1, 3) + dt.timedelta(days=int(r.integers(0, 355)))
            rows.append([b["code"], "Corrective", CM[int(r.integers(0, len(CM)))], d, None, False, None])
    for row in rows:
        u, lag = rq.random(), int(rq.integers(0, 9))
        if u < (.10 if row[1] == "Preventive" else .35) and row[3] < A.CONTRACT_OFF:
            row[3], row[5] = A.workday(row[3]), True
        if row[4] is None:
            row[4] = row[3] - dt.timedelta(days=lag)
    for key in sorted(w.ldar_wos):
        b, day, kind = key
        xs = w.ldar_wos[key]
        raised = max(received(w, x) for x in xs)
        rows.append([b, "Corrective", repair_text(xs, kind), day, min(day, raised) if kind == "routine" else raised,
                     w.ldar["contract"][key], key])
    first = {x["tag"]: x["note"] for x in w.ldar_rows}
    for b, day, point, tag in sorted(w.ldar_rechecks):
        what = fixes(point, FIX)[-1].lower()
        rows.append([b, "Corrective", "Leak repair: %s (%s), %s; %s did not hold" % (
            point, tag, what, first[tag][0].lower() + first[tag][1:]), day, day, False, None])
    rows.sort(key=lambda x: (x[4], x[0], x[2]))
    num = {2025: 25108207, 2026: 26100412, 2027: 27100018}
    for row in rows:
        row.append(num[row[4].year])
        num[row[4].year] += int(rq.integers(1, 6))
    plan = {k: "PM-%05d" % (21003 + 3 * i) for i, k in enumerate(sorted({(x[0], x[2]) for x in rows if x[1] == "Preventive"}))}
    seq, out = {}, []
    for row in sorted(rows, key=lambda x: (x[3], x[0], x[7]) if x[5] else (x[4], x[0], x[7])):
        b, typ, txt, d, raised, con, key, n = row
        if con:
            pre, on = "FT", d
        elif typ == "Preventive":
            pre, on = None, None
        else:
            pre, on = ("EHS" if txt.startswith("Leak repair") else "N"), raised
        if pre:
            seq[(pre, on)] = seq.get((pre, on), 0) + 1
            row.append("%s %s-%02d" % (pre, on.strftime("%y%m%d"), seq[(pre, on)]))
        else:
            row.append(plan[(b, txt)])
        done = A.approval(d) if con else d
        row.append(done)
        if key is not None:
            w.ldar_wo_no[key] = n
    for row in sorted(rows, key=lambda x: (x[9], x[7])):
        b, typ, txt, d, raised, con, key, n, ref, done = row
        if done.year == 2026:
            out.append(dict(WO=n, Battery=code_at(b, d), Type=typ, Description=txt, Reference=ref,
                            Completed=done.isoformat()))
    return write_csv(path, out, list(out[0]))


# ---------------------------------------------------------------- leak detection and repair
LDAR_PULL = dt.date(2027, 1, 12)
# what a maintainer writes for each kind of point; a quick fix first in each list
FIX = [("stuffing box", ["Tightened packing gland", "Adjusted packing", "Replaced packing", "Repacked stuffing box"]),
       ("casing valve", ["Tightened bonnet nuts", "Greased and cycled valve", "Replaced stem packing", "Replaced valve"]),
       ("packing", ["Adjusted packing", "Replaced packing"]),
       ("valve stem", ["Adjusted stem packing", "Repacked stem"]),
       ("flange", ["Re-torqued flange", "Replaced gasket"]),
       ("PSV seat", ["Reseated PSV", "Replaced PSV"]),
       ("sight glass", ["Tightened fitting", "Replaced sight glass gaskets", "Replaced fitting"]),
       ("drain valve", ["Greased and cycled valve", "Replaced valve"]),
       ("threaded", ["Re-taped and made up", "Replaced fitting"]),
       ("orifice", ["Tightened fitting", "Replaced seal ring"]),
       ("burner valve", ["Tightened bonnet nuts", "Replaced valve"]),
       ("regulator", ["Tightened fitting", "Replaced regulator diaphragm"]),
       ("fitting", ["Tightened fitting", "Re-taped and made up", "Replaced fitting"])]
SHUT_FIX = [("stuffing box", ["Repacked stuffing box at shutdown"]), ("flange", ["Replaced gasket at shutdown"]),
            ("valve stem", ["Repacked stem at shutdown"]),
            ("PSV seat", ["Seat lapped at shutdown", "Replaced PSV at shutdown"]),
            ("orifice", ["Replaced seal ring at shutdown"]), ("burner valve", ["Replaced valve at shutdown"]),
            ("packing", ["Repacked at shutdown"]), ("drain valve", ["Replaced valve at shutdown"]),
            ("casing valve", ["Replaced valve at shutdown"]), ("", ["Repaired at shutdown", "Replaced at shutdown"])]


def fixes(point, book):
    return next(xs for key, xs in book if key in point)


def ldar_rows(w):
    """the tracker's rows, one tag for each sighting; a repair's work order closes the tags its maintainer worked, a
    shutdown's the newest tag of each deferred leak"""
    r = np.random.default_rng(W.SEED + 41)
    rows = []
    for i, lk in enumerate(w.ldar["leaks"]):
        closed = lk.get("closed") or lk["fix"]
        done = w.ldar["closure"][(lk["bat"], closed, "shutdown" if lk["deferred"] else "routine")] if closed else None
        for k, (s, rate) in enumerate(zip(lk["seen"], lk["rates"])):
            last = k == len(lk["seen"]) - 1
            fix = closed if last and done is not None and done <= LDAR_PULL else None
            rows.append(dict(leak=i, bat=lk["bat"], point=lk["point"], survey=s, rate=max(1, int(rate)), last=last,
                             deferred=lk["deferred"], fix=fix, done=done if fix else None,
                             recheck=lk["fix"] if "closed" in lk and last else None))
    rows.sort(key=lambda x: (x["survey"], x["bat"], x["point"]))
    nums = r.choice(np.arange(10000, 100000), size=len(rows), replace=False)
    for x, t in zip(rows, nums):
        x["tag"] = "T-%05d" % t
        book = fixes(x["point"], SHUT_FIX if x["deferred"] else FIX)
        u = int(r.integers(0, 12))
        x["note"] = (book[0 if x["recheck"] else u % len(book)] if x["fix"] else
                     ("Needs shutdown" if x["deferred"] else "Awaiting repair"))
    wos = {}
    for x in rows:
        if x["fix"] is not None:
            wos.setdefault((x["bat"], x["fix"], "shutdown" if x["deferred"] else "routine"), []).append(x)
    w.ldar_rows, w.ldar_wos, w.ldar_wo_no = rows, wos, {}
    w.ldar_rechecks = [(x["bat"], x["recheck"], x["point"], x["tag"]) for x in rows if x["recheck"]]


def repair_text(xs, kind):
    what = "; ".join("%s (%s)" % (x["point"], x["tag"]) for x in xs)
    if kind == "shutdown":
        return "Leak repairs at shutdown: %s" % what
    return ("Leak repair: %s" if len(xs) == 1 else "Leak repairs: %s") % what


def ldar_surveys(path, w):
    r = np.random.default_rng(W.SEED + 42)
    tagged = {}
    for x in w.ldar_rows:
        tagged[(x["bat"], x["survey"])] = tagged.get((x["bat"], x["survey"]), 0) + 1
    npts = {}
    for lk in w.ldar["leaks"]:
        npts.setdefault(lk["bat"], set()).add(lk["point"])
    rows = []
    for code in sorted(w.ldar["surveys"]):
        b = w.by[code]
        base = 70 + 30 * sum(1 for x in w.wells if x["bat"] == code) + (60 if b["cls"] != "mwf" else 0) + \
            (90 if b["cls"] == "twice" else 0)
        for d in sorted(w.ldar["surveys"][code]):
            rows.append([d, code_at(code, d), base + int(r.integers(-12, 13)), tagged.get((code, d), 0),
                         w.ldar_crew[(code, d)]])
    # the crew's re-checks of repaired points: each repaired point and the fittings beside it, nothing tagged
    pts_of = {}
    for lk in w.ldar["leaks"]:
        closed = lk.get("closed") or lk["fix"]
        if closed is not None and not lk["deferred"]:
            pts_of.setdefault((lk["bat"], closed, "routine"), set()).add(lk["point"])
    rs = np.random.default_rng(W.SEED + 46)
    for (code, d), ks in sorted(w.ldar["sessions"].items()):
        pts = sorted(set().union(*(pts_of.get(k, set()) for k in ks)))
        last = max(s for s in w.ldar["surveys"][code] if s < d)
        rows.append([d, code_at(code, d), sum(2 + int(rs.integers(0, 3)) for _ in pts), 0, w.ldar_crew[(code, last)]])
    rows.sort(key=lambda x: (x[0], x[1]))
    out = []
    for i, (d, code, pts, n, crew) in enumerate(rows, 1):
        out.append(dict(Survey="LS%02d-%04d" % (d.year % 100, i), Battery=code, Surveyed=d.isoformat(), Technician=crew,
                        Camera="GF320 #%d" % (1 if crew in ("ogi03",) else 2), PointsScreened=pts, LeaksTagged=n))
    return write_csv(path, out, list(out[0]))


def crews(w):
    """two crews share each route's batteries in a round; in 2026 the second crew's tablet syncs its tags to the
    tracker at the end of its week, so the tracker dates them on the Friday"""
    out = {}
    route = {b["code"]: b["area"] for b in w.bats}
    order = {}
    for code in sorted(w.ldar["surveys"]):
        for d in w.ldar["surveys"][code]:
            order.setdefault((route[code], d.year, d.month < 8), []).append((d, code))
    for k, xs in order.items():
        names = ("ogi03", "ogi06") if k[1] == 2025 else ("ogi03", "ogi08")
        for i, (d, code) in enumerate(sorted(xs)):
            out[(code, d)] = names[i % 2]
    w.ldar_crew = out


def synced(d):
    """the Friday a tablet's week of tags reaches the tracker"""
    while d.weekday() != 4:
        d += dt.timedelta(days=1)
    return d


def ldar_tracker(path, w):
    rows = []
    for x in w.ldar_rows:
        key = (x["bat"], x["fix"], "shutdown" if x["deferred"] else "routine") if x["fix"] else None
        wo = w.ldar_wo_no[key] if key else None
        status = "Closed" if x["fix"] else ("Deferred" if x["deferred"] else "Open")
        rows.append([x["tag"], code_at(x["bat"], x["survey"]), received(w, x), x["point"], x["rate"], status,
                     wo, x["done"], x["note"]])
    hdr = ["Tag", "Battery", "Received", "Component", "CH4 rate (g/h)", "Status", "Work order", "Closed", "Note"]
    write_xlsx(path, [("Leak tags", hdr, rows)], widths={"Component": 36, "Note": 28, "Received": 11, "Closed": 11,
                                                          "Work order": 11, "Status": 10, "CH4 rate (g/h)": 13})
    wb = openpyxl.load_workbook(path)
    ws = wb["Leak tags"]
    for row in ws.iter_rows(min_row=2):
        row[2].number_format = "yyyy-mm-dd"
        row[7].number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:I%d" % (len(rows) + 1)
    wb.save(path)
    return len(rows)


# ---------------------------------------------------------------- the two files the return does not draw on
def sales_gas(path, w):
    """the sales meters' monthly statements: the conserved gas the field's three gathering systems delivered"""
    r = np.random.default_rng(W.SEED + 44)
    meters = {"KES-SM-01 Kesmere North": (("R01", "R02", "R03", "R04"), 88.6),
              "KES-SM-02 Kesmere Central": (("R05", "R06", "R07", "R08"), 87.4),
              "KES-SM-03 Kesmere South": (("R09", "R10", "R11", "R12"), 86.9)}
    rows = []
    drift = {name: 0.0 for name in meters}
    for m in range(1, 13):
        for name, (areas, c1) in meters.items():
            vol = sum(p["gas"] for b in w.bats if b["area"] in areas for p in w.prod[b["code"]]
                      if p["day"].year == 2026 and p["day"].month == m) * r.uniform(.9, .94)
            drift[name] = .7 * drift[name] + float(r.normal(0, .22))
            c = c1 + .45 * np.cos(2 * np.pi * (m - 1) / 12) + drift[name]
            hv = 36.95 + .29 * (100 - c) + float(r.normal(0, .04))
            rows.append(dict(Meter=name, Month="2026-%02d" % m, Volume_e3m3="%.1f" % vol, HHV_MJ_m3="%.2f" % hv,
                             Energy_GJ="%.0f" % (vol * hv), C1_molpct="%.2f" % c))
    return write_csv(path, rows, list(rows[0]))


def hauler(path, w):
    r = rng()
    rows = []
    n = 771204
    for b in w.bats:
        acc = 0.0
        for p in w.prod[b["code"]]:
            if p["day"].year != 2026:
                continue
            acc += p["oil"]
            if acc >= 28 + r.uniform(0, 6):
                t = dt.datetime.combine(p["day"], dt.time()) + dt.timedelta(hours=float(r.uniform(7, 19)))
                rows.append([t, code_at(b["code"], p["day"]), round(acc * r.uniform(.985, 1.0), 2),
                             round(float(r.uniform(.2, 1.4)), 1)])
                acc = 0.0
    rows.sort()
    out = []
    for t, code, vol, bsw in rows:
        out.append(dict(Ticket="KT%06d" % n, Battery=code, Loaded=t.strftime("%Y-%m-%d %H:%M"), NetOil_m3="%.2f" % vol,
                        BSW_pct="%.1f" % bsw, Carrier="Kesmere Fluid Haulers"))
        n += 1
    return write_csv(path, out, list(out[0]))


def cut(paths, w):
    """cut every file; paths maps a key to a target path"""
    apply_rename(w)
    w.alarms_by_lift = {}
    for x in w.alarms:
        w.alarms_by_lift.setdefault((x["bat"], x["lift"]), []).append(x)
    vr = visit_rows(w)
    sites = site_ids(w)
    n = {}
    n["register"] = register(paths["register"], w)
    n["fieldlog"] = fieldlog(paths["fieldlog"], w, vr)
    n["sitelink"] = sitelink(paths["sitelink"], w, vr, sites)
    n["sites"] = sites_csv(paths["sites"], w, sites)
    n["events"] = event_register(paths["events"], w)
    n["alarms"] = alarm_journal(paths["alarms"], w)
    n["production"] = production(paths["production"], w)
    n["vapour"] = vapour_analyses(paths["vapour"], w)
    ldar_rows(w)
    crews(w)
    n["workorders"] = work_orders(paths["workorders"], w)
    n["ldar_surveys"] = ldar_surveys(paths["ldar_surveys"], w)
    n["ldar_tags"] = ldar_tracker(paths["ldar_tags"], w)
    n["hauler"] = hauler(paths["hauler"], w)
    n["salesgas"] = sales_gas(paths["salesgas"], w)
    w.sites = sites
    return n
