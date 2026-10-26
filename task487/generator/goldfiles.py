"""The golden deliverables from the cut pack: the Board's event schema for the 2026 return (release_events_2026.csv) and
the regulatory analyst's return workbook (the year's methane, its band and the 2027 levy, and Schedule B: each route's
fugitive methane in each half of the year). Every figure is computed once here and returned for
figures.json."""
import csv
import datetime as dt

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import family as FAM
import golden as G
import pack_read as PR

KIND_NAME = {"TH": "Open thief hatch", "CP": "Combustor pilot out", "PV": "Relief valve not reseated"}
ROUTES = ["R%02d" % i for i in range(1, 13)]


def compute(p=None):
    p = p or PR.read()
    g = G.compute(p)
    rows = g["rows"]
    fam = FAM.standing(p, g)
    f = {"call": {"band": g["band"], "levy_cad": g["levy"], "rate_per_battery_cad": g["rate"], "batteries": g["held"]}}
    f["methane_t"] = round(g["total"], 1)
    f["methane_t.family_lo"] = round(min(fam.values()), 1)
    f["methane_t.family_hi"] = round(max(fam.values()), 1)
    f["events.total"] = len(rows)
    for k, nm in (("TH", "hatch"), ("CP", "pilot"), ("PV", "relief")):
        f["events.%s" % nm] = sum(1 for r in rows if r["kind"] == k)
    f["events.relief_alarm_timed"] = sum(1 for r in rows if r["basis"] == "ALARM")
    for rt in ROUTES:
        for h in ("h1", "h2"):
            f["schedule_b.%s.%s_t" % (rt, h)] = round(g["sched"][rt][h], 1)
    reg, cur = G.codes(p)
    h1, h2, _, _ = G.runs(p, cur)
    f["schedule_b.moved_jul"] = sum(1 for b in h1 if b in h2 and h2[b] != h1[b])
    f["schedule_b.moved_jan"] = sum(1 for b, r in reg.items() if b in h2 and r["route"] != h2[b])
    wos = G.leak_wos(p)
    f["schedule_b.rechecked_tags"] = sum(1 for r in p["ldar_tags"] if wos.get(r["Tag"]) and (
        G.as_date(r["Closed"]) is None or max(wos[r["Tag"]]) > G.as_date(r["Closed"])))
    sv = G.surveys(p, g["opts"], cur)
    f["schedule_b.sessions"] = sum(1 for r in p["ldar_surveys"] if dt.date.fromisoformat(r["Surveyed"]) not in
                                   sv[cur.get(r["Battery"], r["Battery"])])
    f["schedule_b.contract_orders"] = sum(1 for r in p["workorders"] if r["Description"].startswith("Leak repair") and
                                          r["Reference"].startswith("FT "))
    return p, g, f, fam


def local_iso(h):
    return G.local_dt(h).strftime("%Y-%m-%dT%H:%M")


def events_csv(path, g):
    rows = sorted(g["rows"], key=lambda r: (r["found"], r["no"]))
    with open(path, "w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh, lineterminator="\n")
        wr.writerow(["event_ref", "facility_licence", "battery", "kind", "start", "start_basis", "found", "closed",
                     "active_hours", "methane_kg"])
        for r in rows:
            wr.writerow([r["no"], r["licence"], r["bat"], G.KIND_CODE[r["kind"]], local_iso(r["start"]), r["basis"],
                         local_iso(r["found"]), local_iso(r["closed"]), "%.2f" % r["hours"], "%.1f" % r["kg"]])
    return len(rows)


def _w(ws, widths):
    for i, x in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


def _head(row, fill, bold):
    for c in row:
        c.font = bold
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="top")


def workbook(path, p, g, f):
    bold = Font(bold=True)
    head = PatternFill("solid", fgColor="E3E9EF")
    key = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="9AA5B1")
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Return"
    ws.append(["2026 release-event return: %.1f t of methane, band %s, a 2027 levy of CAD %s" % (
        g["total"], g["band"], format(g["levy"], ","))])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Vennholt Energy Ltd., Kesmere field, under Directive 17 (2025 edition). For the area production "
               "superintendent's sign-off; due to the Provincial Energy Board 31 March 2027."])
    ws.append([])
    ws.append(["Item", "Value", "Basis"])
    _head(ws[4], head, bold)
    lines = [("Release-event methane, 2026 (t)", g["total"], "Directive 17 s.3; events found in 2026"),
             ("Band", g["band"], "Directive 17 Table 1 (50 t to under 60 t)"),
             ("2027 levy rate, each battery (CAD)", g["rate"], "Directive 17 Table 1"),
             ("Batteries held at the start of 2027", g["held"], "battery_register_20270108.csv"),
             ("2027 levy (CAD)", g["levy"], "rate times batteries held (s.4)"),
             ("Release events found in 2026", f["events.total"], "release_events_2026.csv")]
    for nm, v, basis in lines:
        ws.append([nm, v, basis])
    ws["B5"].number_format = "0.0"
    ws["B7"].number_format = ws["B9"].number_format = "#,##0"
    for r in (5, 6, 9):
        ws.cell(r, 1).fill = ws.cell(r, 2).fill = key
        ws.cell(r, 2).font = bold
    ws.append([])
    ws.append(["Events by kind", "Events", "Start basis"])
    _head(ws[ws.max_row], head, bold)
    for k in ("TH", "CP", "PV"):
        n = sum(1 for r in g["rows"] if r["kind"] == k)
        basis = "expected start" if k != "PV" else "alarm at telemetered batteries (%d), expected start elsewhere (%d)" % (
            f["events.relief_alarm_timed"], n - f["events.relief_alarm_timed"])
        ws.append([KIND_NAME[k], n, basis])
    ws.append(["Total", f["events.total"], None])
    for c in ws[ws.max_row]:
        c.font = bold
        c.border = Border(top=thin)
    ws.freeze_panes = "A5"
    _w(ws, [40, 14, 66])

    s = wb.create_sheet("Schedule B")
    s.append(["Schedule B: fugitive methane by route, 2026 (tonnes; not levied), each half from the batteries the "
              "route ran then"])
    s["A1"].font = Font(bold=True, size=12)
    s.append([])
    s.append(["Route", "January to June", "July to December", "Year"])
    _head(s[3], head, bold)
    for rt in ROUTES:
        v = g["sched"][rt]
        s.append([rt, v["h1"], v["h2"], v["h1"] + v["h2"]])
    fld = g["sched"]["field"]
    s.append(["Total", fld["h1"], fld["h2"], fld["h1"] + fld["h2"]])
    for row in s.iter_rows(min_row=4, max_row=s.max_row):
        for c in row[1:]:
            c.number_format = "#,##0.0"
    for c in s[s.max_row]:
        c.font = bold
        c.border = Border(top=thin)
    s.freeze_panes = "B4"
    _w(s, [12, 17, 17, 12])

    n = wb.create_sheet("Notes")
    notes = [
        "Release events are those found at the Kesmere batteries in calendar 2026 (Directive 17 s.3.1), from the "
        "office register, each joined to the lease operator visit that found it in the route app (to 30 June) or the check-in app "
        "(from 1 July; R03 and R09 from 15 June, logged in both apps to 30 June and counted once).",
        "A relief valve that did not reseat at a telemetered battery starts at the first alarm of the last lift the "
        "SCADA journal records before its find, on the RTU's own clock.",
        "Every other start is its expected value given the records (s.3.2): for each kind, the chance that a window "
        "ends with an event of that kind is fitted on all of the operator's 2026 windows as a share started as the "
        "previous visit left plus a rate an hour between visits, and each event's start is its expected value given "
        "its window.",
        "Methane in each hour is the battery's tank vapour methane that day (s.3.3): the day's oil, the flash gas-oil "
        "ratio and methane mole fraction of the latest tank vapour analysis on or before the day, 0.6785 kg a cubic "
        "metre, over 24 hours. BT-0907A carries BT-0907's records to 16 August 2026.",
        "Fugitive emissions (s.5): a battery's surveys are the camera sessions that screened its points; the crews' "
        "re-checks of a few repaired fittings in the same log are not surveys. Each tag belongs to its battery's last "
        "survey on or before the day the tag reached the tracker; a tag hung on a component whose earlier tag had not "
        "been repaired by that survey is the same leak found again. A leak is repaired on the day of the last leak "
        "repair that names its tag: a contract work order on its field ticket's day (the maintenance system completes "
        "a contract order when the foreman approves its tickets, on the 10th of the next month), and a tag closed on "
        "a first repair that a re-check found still leaking runs to the later repair. Each leak vents from its "
        "battery's last survey before it was first found to its repair, at the rate measured then, split at 1 July "
        "and kept inside 2026.",
        "Each half's fugitive methane is carried by the run that served the battery in that half: the route app's "
        "Route to 30 June, the check-in app's run from 1 July. The battery register carries the runs redrawn for "
        "2027.",
        "Figures are computed unrounded and shown rounded; totals are of unrounded values."]
    n.append(["Notes"])
    n["A1"].font = Font(bold=True, size=12)
    for t in notes:
        n.append([t])
        n.cell(n.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    _w(n, [120])
    wb.save(path)


def write(gpath):
    p, g, f, fam = compute()
    events_csv(gpath("events"), g)
    workbook(gpath("workbook"), p, g, f)
    return f
