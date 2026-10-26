"""Read the cut pack back into rows: golden.py, ceiling.py and checks.py compute from the pack as shipped, never from
the world."""
import csv
import datetime as dt
import json
import os

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
TARGET = os.path.join(os.path.dirname(HERE), "target")
REGISTER = "battery_register_20270108.csv"
FIELDLOG = "route_app_visits_20251215_20260630.csv"
SITELINK = "checkin_app_visits_20260615_20270112.jsonl"
SITES = "checkin_app_sites.csv"
EVENTS = "release_event_register_2026.xlsx"
ALARMS = "scada_alarm_journal_20251215_20270112.csv"
PRODUCTION = "battery_daily_production_20251215_20261231.csv"
VAPOUR = "tank_vapour_analyses_2025_2026.xlsx"
WORKORDERS = "maintenance_work_orders_2026.csv"
LDAR_SURVEYS = "ldar_surveys_2025_2026.csv"
LDAR_TAGS = "ldar_leak_tracker_20270112.xlsx"
DIRECTIVE = "peb_directive_17_release_event_return_2025.pdf"


def rows(name, target=TARGET):
    return list(csv.DictReader(open(os.path.join(target, name), encoding="utf-8")))


def sheet(name, target=TARGET):
    wb = openpyxl.load_workbook(os.path.join(target, name), read_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    hdr = next(it)
    return [dict(zip(hdr, r)) for r in it]


def jsonl(name, target=TARGET):
    return [json.loads(l) for l in open(os.path.join(target, name), encoding="utf-8") if l.strip()]


def read(target=TARGET):
    return dict(register=rows(REGISTER, target), fieldlog=rows(FIELDLOG, target), sitelink=jsonl(SITELINK, target),
                sites=rows(SITES, target), events=sheet(EVENTS, target), alarms=rows(ALARMS, target),
                production=rows(PRODUCTION, target), vapour=sheet(VAPOUR, target),
                workorders=rows(WORKORDERS, target), ldar_surveys=rows(LDAR_SURVEYS, target),
                ldar_tags=sheet(LDAR_TAGS, target))
