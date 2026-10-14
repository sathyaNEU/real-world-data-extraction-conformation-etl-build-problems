"""The goldens, computed from the shipped pack alone: the determination for the board (PDF, with the chart of every
offered plant's rate against the standard) and the workbook of every offered plant; generator/figures.json carries
every figure the write-up states."""
import io
import json
import math
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import calc as C        # noqa: E402
import world as W       # noqa: E402

PDF = "eps_determination_rfo_2027.pdf"
XLSX = "offered_plants_eps_2027.xlsx"
SHORT = {"Sorrel Butte CCGT": "Sorrel Butte", "Kessling Creek Generating Station": "Kessling Creek",
         "Tamarack Landing": "Tamarack Landing", "Ochre Flat 7HA": "Ochre Flat",
         "Brenmoor Mill Combined Cycle": "Brenmoor Mill", "Juniper Draw": "Juniper Draw",
         "Lindgren Hill Generating Station": "Lindgren Hill", "Ashby Gulch CC": "Ashby Gulch",
         "Moraine Point": "Moraine Point", "Coldbrook Combined Cycle": "Coldbrook"}


def compute(target):
    L = C.load(target)
    R = C.routes(L)
    FO = C.forced_outages(target)
    SH = C.substitute_hours(L)
    rows = {}
    for oid in sorted(L["cap"]):
        r = R[oid]
        fam = [r["family"][k] for k in C.GOLDEN]
        lo, hi = min(fam), max(fam)
        comp = r["golden"] <= C.STANDARD
        assert all((v <= C.STANDARD) == comp for v in fam), (oid, fam)
        rows[oid] = dict(name=SHORT[L["name"][oid]], resource=L["name"][oid], cap=L["cap"][oid], rate=r["golden"],
                         lo=int(math.floor(lo)), hi=int(math.ceil(hi)), n=r["n"], complies=comp,
                         default=r["default"], no_starts=r["no_starts"], top=r["top"],
                         fo_hours=FO[oid][0], fo_count=FO[oid][1], sub_hours=SH[oid])
    offers = {o["Offer ID"]: o for o in L["offers"]}
    excess = {o: C.low_output_excess(L, o) for o in rows}
    complying = [o for o in sorted(rows) if rows[o]["complies"]]
    taken, total, passed = C.slate(L, complying)
    trace, run = [], 0.0
    for o in sorted(complying, key=lambda o: (L["price"][o], o)):
        if o in taken:
            run += rows[o]["cap"]
            trace.append((o, True, run))
        else:
            trace.append((o, False, run + rows[o]["cap"]))
    dflt = [o for o in sorted(rows) if rows[o]["default"] <= C.STANDARD]
    top = [o for o in sorted(rows) if rows[o]["top"] <= C.STANDARD]
    unf = complying + [o for o in rows if rows[o]["name"] == "Tamarack Landing"]
    return dict(L=L, R=R, rows=rows, offers=offers, complying=complying, signed=taken, mw=total, passed=passed,
                trace=trace, excess=excess, offered=sum(r["cap"] for r in rows.values()),
                complying_mw=sum(rows[o]["cap"] for o in complying),
                default_complying=dflt, default_slate=C.slate(L, dflt), top_complying=top, top_slate=C.slate(L, top),
                unfired_slate=C.slate(L, unf))


def names_list(names):
    return ", ".join(names[:-1]) + " and " + names[-1] if len(names) > 1 else names[0]


def chart_png(X):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams["font.family"] = "DejaVu Sans"
    rows = X["rows"]
    order = sorted(rows, key=lambda o: rows[o]["rate"])
    vals = [rows[o]["rate"] for o in order]
    labels = ["%s (%d MW)" % (rows[o]["name"], rows[o]["cap"]) for o in order]
    signed = set(X["signed"])
    cols = ["#1f4e5f" if o in signed else ("#7fa7b5" if rows[o]["complies"] else "#c9a27e") for o in order]
    fig, ax = plt.subplots(figsize=(6.4, 3.3), dpi=150)
    ax.barh(range(len(order)), vals, color=cols, height=0.62)
    ax.axvline(C.STANDARD, color="#a93226", lw=1.3, ls="--")
    ax.text(C.STANDARD + 4, len(order) - 0.45, "standard, 925", color="#a93226", fontsize=7.5, va="center")
    for k, v in enumerate(vals):
        ax.text(v + 5, k, "{:,}".format(int(round(v))), va="center", ha="left", fontsize=7.3)
    ax.set_yticks(range(len(order)))
    ax.set_yticklabels(labels, fontsize=7.8)
    ax.set_xlim(700, max(vals) + 70)
    ax.set_xlabel("Pounds of CO2 per net MWh at the Contract Capacity, 2024 and 2025 records", fontsize=7.8)
    ax.tick_params(axis="x", labelsize=7.5)
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    from matplotlib.patches import Patch
    ax.legend(handles=[Patch(color="#1f4e5f", label="signed"), Patch(color="#7fa7b5", label="complies, passed over"),
                       Patch(color="#c9a27e", label="above the standard")], fontsize=7, frameon=False,
              loc="lower right")
    ax.set_title("%s offered plants meet the standard; %s of them fill the need" % (
        ("Seven" if len(X["complying"]) == 7 else len(X["complying"])), ("three" if len(signed) == 3 else len(signed))),
        fontsize=9.5, loc="left")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def write_pdf(X, path):
    from reportlab import rl_config
    rl_config.invariant = 1
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    FD = "/usr/share/fonts/truetype/liberation/"
    for nm, f in (("LiberationSans", "Regular"), ("LiberationSans-Bold", "Bold"), ("LiberationSans-Italic", "Italic")):
        pdfmetrics.registerFont(TTFont(nm, FD + nm.split("-")[0] + "-" + f + ".ttf"))
    pdfmetrics.registerFontFamily("LiberationSans", normal="LiberationSans", bold="LiberationSans-Bold",
                                  italic="LiberationSans-Italic", boldItalic="LiberationSans-Bold")
    INK, GREY = colors.HexColor("#1b1d20"), colors.HexColor("#5b6068")
    B = ParagraphStyle("b", fontName="LiberationSans", fontSize=9.6, leading=13.4, textColor=INK, spaceAfter=6)
    H = ParagraphStyle("h", fontName="LiberationSans-Bold", fontSize=12.5, leading=16, textColor=INK, spaceAfter=4)
    H2 = ParagraphStyle("h2", fontName="LiberationSans-Bold", fontSize=10, leading=13, textColor=INK, spaceBefore=6, spaceAfter=3)
    SM = ParagraphStyle("sm", fontName="LiberationSans", fontSize=7.8, leading=10, textColor=GREY, spaceAfter=3)
    CE = ParagraphStyle("ce", fontName="LiberationSans", fontSize=8.4, leading=10.5, textColor=INK)
    CR = ParagraphStyle("cr", parent=CE, alignment=2)
    CB = ParagraphStyle("cb", parent=CE, fontName="LiberationSans-Bold")
    CBR = ParagraphStyle("cbr", parent=CB, alignment=2)
    gm, gmt = W.PEOPLE["gm"]
    pl, plt_ = W.PEOPLE["planning"]
    L = X["L"]
    rows = X["rows"]
    signed = X["signed"]
    out = [o for o in sorted(rows) if not rows[o]["complies"]]
    signed_names = ["%s (%d MW at $%.2f%s)" % (rows[o]["name"], rows[o]["cap"], L["price"][o],
                                               " per kW-month" if k == 0 else "") for k, o in enumerate(signed)]
    first_pass = X["passed"][0]
    before = [t for t in X["trace"] if t[0] == first_pass][0]

    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFont("LiberationSans", 7.2)
        canvas.setFillColor(GREY)
        canvas.drawString(0.9 * inch, 0.55 * inch, "%s  |  Emissions performance determination, %s  |  6 November 2026"
                          % (W.DISTRICT, W.RFO))
        canvas.drawRightString(LETTER[0] - 0.9 * inch, 0.55 * inch, "Page %d" % doc.page)
        canvas.restoreState()

    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=0.9 * inch, rightMargin=0.9 * inch, topMargin=0.75 * inch,
                            bottomMargin=0.85 * inch, title="Emissions performance determination, 2026 Baseload Request for Offers",
                            author=W.DISTRICT, subject="", creator="", producer="")
    s = [Paragraph(W.DISTRICT.upper(), SM),
         Paragraph("Sign %s: %s MW of the %s MW baseload need" % (names_list([rows[o]["name"] for o in signed]),
                                                                  "{:,}".format(int(X["mw"])), "{:,}".format(int(W.NEED))), H),
         Paragraph("To: Board of Commissioners<br/>From: %s, %s<br/>Through: %s, %s<br/>Date: 6 November 2026<br/>"
                   "Re: Emissions performance determination under RP-7 Section 5, %s" % (pl, plt_, gm, gmt, W.RFO), B),
         Spacer(1, 4),
         Paragraph("%s of the ten plants offered emit at or under the state standard of 925 pounds of carbon dioxide per "
                   "net megawatt-hour under their offered commitments. Taken in order of capacity price, each offer "
                   "whole, passing over any that would take the District above its 1,000 MW need (request for offers "
                   "5.2), they fill %s MW: %s. %s, second by price, is passed over: with %s it would make %s MW. %s "
                   "emit more than 925 pounds under their offers and cannot be signed."
                   % (("Seven" if len(X["complying"]) == 7 else str(len(X["complying"]))), "{:,}".format(int(X["mw"])),
                      names_list(signed_names), rows[first_pass]["name"], rows[signed[0]]["name"],
                      "{:,}".format(int(before[2])), names_list([rows[o]["name"] for o in out])), B)]
    head = [Paragraph(t, CB if k < 2 else CBR) for k, t in enumerate(
        ("Offer", "Plant", "Contract Capacity (MW)", "Capacity price ($/kW-mo)", "lb CO2 per net MWh", "Complies", "Signed"))]
    data = [head]
    for o in sorted(rows):
        r = rows[o]
        data.append([Paragraph(o, CE), Paragraph(r["name"], CE), Paragraph("%d" % r["cap"], CR),
                     Paragraph("%.2f" % L["price"][o], CR), Paragraph("{:,}".format(int(round(r["rate"]))), CR),
                     Paragraph("Yes" if r["complies"] else "No", CR), Paragraph("Yes" if o in signed else "", CR)])
    t = Table(data, colWidths=[0.7 * inch, 1.45 * inch, 1.05 * inch, 1.05 * inch, 1.0 * inch, 0.75 * inch, 0.6 * inch],
              hAlign="LEFT")
    t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, 0), 0.6, INK), ("LINEBELOW", (0, -1), (-1, -1), 0.4, GREY),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 1.5),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5)]))
    s.append(t)
    s.append(Paragraph("Source: EPA Clean Air Markets Program Data, hourly emissions for 2024 and 2025 (downloaded 6 October "
                       "2026); respondents' revenue-meter data; offers received to 30 September 2026; District analysis. "
                       "Rates are compared with the standard unrounded (RP-7 5.4) and shown to the nearest pound.", SM))
    s.append(Paragraph("Kessling Creek and Ashby Gulch left the EIA plant code blank on their offers; each was matched to "
                       "its plant by name and county in Form EIA-860. Moraine Point's monitored unit reports as CT-1 to "
                       "30 June 2025 and as 1A after its data acquisition system was replaced; both are counted.", SM))
    s.append(Paragraph("How the need was filled", H2))
    steps = []
    for o, took, run in X["trace"]:
        steps.append("%s, %d MW at $%.2f: %s" % (rows[o]["name"], rows[o]["cap"], L["price"][o],
                                                 ("taken, %s MW" % "{:,}".format(int(run))) if took else
                                                 ("passed over, it would make %s MW" % "{:,}".format(int(run)))))
    s.append(Paragraph("The complying offers in order of capacity price: " + "; ".join(steps) + ".", B))
    s.append(Spacer(1, 4))
    s.append(Image(io.BytesIO(chart_png(X)), width=5.6 * inch, height=2.89 * inch))
    s.append(Paragraph("Figure 1. Each offered plant's carbon dioxide per net megawatt-hour at its Contract Capacity "
                       "against the 925-pound standard; the plants signed darkest, the complying offers passed over "
                       "lighter.", SM))
    s.append(Paragraph("How the rates were determined", H2))
    ns = [rows[o]["n"] for o in rows]
    s.append(Paragraph(
        "Each agreement has the plant deliver its Contract Capacity in every hour of the delivery term, and the plant may "
        "sell to no one else (request for offers 3.2); RP-7 5.2 rates a plant on the carbon dioxide it emits per net "
        "megawatt-hour it delivers under the commitment. A combined cycle's carbon dioxide per megawatt-hour depends on "
        "the output it runs at: part load and the hours after a start burn more fuel for each megawatt-hour, and duct "
        "burners add output at a far higher heat rate than the turbines. Each plant's rate was therefore read at its "
        "Contract Capacity, as EPA's hourly carbon dioxide over the revenue meter's net generation in the 2024 and 2025 "
        "hours in which the meter shows the plant delivering its Contract Capacity (within 2 per cent, every monitored "
        "unit operating the full hour): %d to %s hours a plant. A window of 1, 3 or 5 per cent, 2025 alone, or a curve "
        "fitted to each plant's hourly records leaves every plant on the same side of the standard." % (
            min(ns), "{:,}".format(max(ns))), B))
    cyc = [o for o in sorted(rows) if rows[o]["complies"] and rows[o]["default"] > C.STANDARD]
    fired = [o for o in sorted(rows) if not rows[o]["complies"] and rows[o]["default"] <= C.STANDARD]
    tl = rows[fired[0]]
    ds, dmw, _ = X["default_slate"]
    s.append(Paragraph(
        "The plants' 2025 annual averages would give a different list. %s cycled through 2025, starting most days and "
        "running much of the time at part load, and averaged %s to %s pounds for the year; at their Contract "
        "Capacities they emit %s to %s. %s averaged %s pounds in 2025, running mostly at its unfired base load, but its "
        "%d MW Contract Capacity is more than its turbines make without the duct burners in any season, so it fires "
        "them in every contract hour and emits %s. On the averages the fill would take %s, %s MW." % (
            names_list([rows[o]["name"] for o in cyc]), int(round(min(rows[o]["default"] for o in cyc))),
            int(round(max(rows[o]["default"] for o in cyc))), int(round(min(rows[o]["rate"] for o in cyc))),
            int(round(max(rows[o]["rate"] for o in cyc))), tl["name"], int(round(tl["default"])), tl["cap"],
            int(round(tl["rate"])), names_list([rows[o]["name"] for o in ds]), "{:,}".format(int(dmw))), B))
    s.append(Paragraph("Availability and the records behind the rates", H2))
    fo = sorted(rows, key=lambda o: -rows[o]["fo_hours"])
    s.append(Paragraph(
        "The workbook also gives each plant's 2025 forced outages from its GADS records, counting the hours the whole "
        "plant was out (every combustion turbine out on a forced event; a single turbine or a steam turbine the "
        "turbines can bypass leaves the plant running) from each unit's latest entry for an outage, since three first "
        "estimates of return to service ran past hours in which EPA's records show the plant running, and the 2025 "
        "hours in which a monitored unit's heat input, and "
        "with it the carbon dioxide EPA calculates from it, came from substitute data. %s had the most forced outage "
        "hours (%d in %d outages)." % (rows[fo[0]]["name"], int(round(rows[fo[0]]["fo_hours"])), rows[fo[0]]["fo_count"]), B))
    s.append(Spacer(1, 6))
    s.append(Paragraph("%s<br/>%s" % (pl, plt_), B))
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


def write_xlsx(X, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    rows = X["rows"]
    L = X["L"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Offered plants"
    head = ["Offer ID", "Plant", "Respondent", "Contract Capacity (MW)", "Capacity price ($/kW-month, 2027)",
            "CO2 rate under the commitment (lb per net MWh)", "2025 forced outage hours", "2025 forced outages",
            "2025 clock hours operated with CO2 resting on substitute data"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for o in sorted(rows):
        r = rows[o]
        ws.append([o, r["name"], X["offers"][o]["Respondent"], int(r["cap"]), L["price"][o], int(round(r["rate"])),
                   int(round(r["fo_hours"])), r["fo_count"], r["sub_hours"]])
    ws.append([None, "All offers", None, int(X["offered"]), None, None, None, None, None])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for col, wdt in zip("ABCDEFGHI", (10, 18, 32, 14, 16, 20, 14, 12, 20)):
        ws.column_dimensions[col].width = wdt
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            if isinstance(c.value, int):
                c.number_format = "#,##0"
            elif isinstance(c.value, float):
                c.number_format = "0.00"
    ws.freeze_panes = "C2"
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    n = wb.create_sheet("Notes")
    for line in (
            ["Prepared by Resource Planning, 6 November 2026, from the evaluation folder of 4 November 2026 (offers "
             "received to 30 September 2026; EPA data downloaded 6 October 2026)."],
            ["CO2 rate under the commitment: EPA hourly CO2 (40 CFR 75, substitute data as reported) over revenue-meter "
             "net generation in the 2024 and 2025 hours the plant delivered its Contract Capacity (within 2 per cent, "
             "full operating hours). Compared with 925 unrounded (RP-7 5.4)."],
            ["Forced outage hours and forced outages: hours in 2025 (Pacific time) the whole plant was off line on a "
             "forced event (GADS types U1, U2, U3, SF), which is when every combustion turbine, or the single-shaft "
             "unit, was out; each unit's latest entry for an outage's start, merged in time; Moraine Point's outage from "
             "New Year's Eve counted for its 2025 hours."],
            ["Clock hours operated with CO2 resting on substitute data: 2025 clock hours (local standard time) in which a "
             "monitored unit operated with its heat input reported from substitute data, from which EPA's CO2 is "
             "calculated (Appendix G); an hour counts once however many units it covers."]):
        n.append(line)
    n.column_dimensions["A"].width = 120
    wb.save(path)


def figures(X):
    rows = X["rows"]
    L = X["L"]
    ds, dmw, dps = X["default_slate"]
    ts, tmw, _ = X["top_slate"]
    us, umw, _ = X["unfired_slate"]
    first = X["passed"][0]
    F = {"slate": [rows[o]["name"] for o in X["signed"]], "slate_mw": int(X["mw"]), "need_mw": int(W.NEED),
         "offered_mw": int(X["offered"]), "standard": 925, "complying": [rows[o]["name"] for o in X["complying"]],
         "complying_count": len(X["complying"]), "complying_mw": int(X["complying_mw"]),
         "not_complying": [rows[o]["name"] for o in sorted(rows) if not rows[o]["complies"]],
         "first_passed_over": rows[first]["name"], "first_passed_over_total_mw": int([t for t in X["trace"] if t[0] == first][0][2]),
         "fill_trace": [[rows[o]["name"], took, int(run)] for o, took, run in X["trace"]],
         "default_complying": [rows[o]["name"] for o in X["default_complying"]],
         "default_slate": [rows[o]["name"] for o in ds], "default_slate_mw": int(dmw),
         "top_complying": [rows[o]["name"] for o in X["top_complying"]],
         "top_slate": [rows[o]["name"] for o in ts], "top_slate_mw": int(tmw),
         "unfired_slate": [rows[o]["name"] for o in us], "unfired_slate_mw": int(umw),
         "hours_min": min(r["n"] for r in rows.values()), "hours_max": max(r["n"] for r in rows.values()),
         "window_per_cent": 2, "variant_windows_per_cent": [1, 3, 5], "curve_from_per_cent": [80, 90],
         "low_output_per_cent": 75, "plants": {}}
    ex = [X["excess"][o] for o in sorted(rows) if X["excess"][o] is not None]
    F["low_output_excess_per_cent"] = [int(round(min(ex))), int(round(max(ex)))]
    F["low_output_plants"] = [rows[o]["name"] for o in sorted(rows) if X["excess"][o] is not None]
    cyc = [o for o in sorted(rows) if rows[o]["complies"] and rows[o]["default"] > C.STANDARD]
    F["cycling"] = [rows[o]["name"] for o in cyc]
    F["cycling_default_range"] = [int(round(min(rows[o]["default"] for o in cyc))), int(round(max(rows[o]["default"] for o in cyc)))]
    F["cycling_golden_range"] = [int(round(min(rows[o]["rate"] for o in cyc))), int(round(max(rows[o]["rate"] for o in cyc)))]
    for o in sorted(rows):
        r = rows[o]
        F["plants"][r["name"]] = {
            "offer_id": o, "contract_mw": int(r["cap"]), "price": "%.2f" % L["price"][o], "rate": int(round(r["rate"])),
            "rate_range": [r["lo"], r["hi"]], "complies": r["complies"], "signed": o in X["signed"], "hours": r["n"],
            "default_2025": int(round(r["default"])), "no_starts_2025": int(round(r["no_starts"])),
            "top_output": int(round(r["top"])), "forced_outage_hours": int(round(r["fo_hours"])),
            "forced_outage_hours_exact": round(r["fo_hours"], 2), "forced_outages": r["fo_count"],
            "substitute_hours": r["sub_hours"]}
    return F


def main(target, out_dir, fig_path):
    os.makedirs(out_dir, exist_ok=True)
    X = compute(target)
    write_pdf(X, os.path.join(out_dir, PDF))
    import postprocess as PP
    PP.normalize_pdf(os.path.join(out_dir, PDF), b"D:20261106091500-08'00'", b"D:20261106102245-08'00'")
    write_xlsx(X, os.path.join(out_dir, XLSX))
    pl = W.PEOPLE["planning"][0]
    PP.normalize_ooxml(os.path.join(out_dir, XLSX), pl, pl, "2026-11-05T22:40:00Z", "2026-11-06T17:58:12Z", excel=True,
                       company=W.DISTRICT)
    F = figures(X)
    with open(fig_path, "w") as fh:
        json.dump(F, fh, indent=1, ensure_ascii=False)
        fh.write("\n")
    for f in (PDF, XLSX):
        os.utime(os.path.join(out_dir, f), (1793989500, 1793989500))
    return X


if __name__ == "__main__":
    T = os.path.dirname(HERE)
    main(os.path.join(T, "target"), os.path.join(T, "golden"), os.path.join(HERE, "figures.json"))
