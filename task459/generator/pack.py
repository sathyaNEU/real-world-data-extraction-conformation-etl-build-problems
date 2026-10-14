"""The pack's data files, written from the hourly truth in sim.py: EPA's hourly emissions for the offered plants'
monitored stacks and their monitoring plans, the EPA-EIA crosswalk, the generator inventory (EIA-860) and the
generation and fuel survey (EIA-923), the offers, and the sellers' offer data packages as the district compiled them
(hourly revenue-meter output, 2025 GADS events, 2027 planned outages)."""
import csv
import datetime as dt
import math
from zoneinfo import ZoneInfo

import numpy as np

import sim as S
import world as W

LA = ZoneInfo("America/Los_Angeles")
UTC = dt.timezone.utc
PST = dt.timezone(dt.timedelta(hours=-8))

FILES = {
    "camd": "camd_hourly_emissions_offered_units_2024-2025.csv",
    "plans": "camd_monitoring_plans_offered_units.csv",
    "xwalk": "epa_eia_crosswalk_2025.csv",
    "eia860": "eia860_2025_generator_offered_plants.xlsx",
    "eia923": "eia923_2024-2025_generation_and_fuel_offered_plants.xlsx",
    "offers": "rfo2026_offers_received.xlsx",
    "meter": "rfo2026_offer_meter_data_2024-2025.csv",
    "gads": "rfo2026_offer_gads_events_2025.csv",
    "outages27": "rfo2026_offer_planned_outages_2027.csv",
    "procedure": "calder_pud_rp7_resource_procurement_procedure.pdf",
    "terms": "rfo2026_baseload_request_for_offers.pdf",
    "dictionary": "rfo2026_evaluation_data_dictionary.txt",
}

# offers in the order the district logged them
OFFER_ORDER = ["TL", "JD", "SB", "CB", "MP", "OF", "LH", "AG", "KC", "PM"]
OFFER_ID = {k: "BL26-%02d" % (i + 1) for i, k in enumerate(OFFER_ORDER)}
RECEIVED = {"TL": "2026-09-22", "JD": "2026-09-24", "SB": "2026-09-25", "CB": "2026-09-28", "MP": "2026-09-29",
            "OF": "2026-09-29", "LH": "2026-09-30", "AG": "2026-09-30", "KC": "2026-09-30", "PM": "2026-09-30"}


def _w(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        wr = csv.writer(f, lineterminator="\r\n")
        wr.writerow(header)
        wr.writerows(rows)
    return len(rows)


def t_of(h):
    return dt.datetime.combine(W.FIRST_DAY, dt.time()) + dt.timedelta(hours=int(h))


def fmt(v, nd):
    return ("%%.%df" % nd) % v


# ------------------------------------------------------------------------------------------------ stack detail
def stack_series(world, key):
    """per monitored stack: operating time, gross load (MW, average over operating time), heat input and the
    substitute flags, from the plant's hourly truth"""
    p = W.P[key]
    r = world["plants"][key]
    N = S.N
    on, u, Ptot, hi = r["on"], r["op"], r["Ptot"], r["hi"]
    mode = r["mode"]
    out = {}
    if p["config"] == "2x1":
        c1 = r["unit_out"]["CT1"] > 0.5
        c2 = r["unit_out"]["CT2"] > 0.5
        half = mode == "half"
        run1 = on & ~(half & c1)
        run2 = on & ~(half & c2)
        nrun = run1.astype(int) + run2.astype(int)
        share = np.where(nrun > 0, 1.0 / np.maximum(nrun, 1), 0.0)
        ct_part = 0.66 if key in W.TURBINE_ONLY_LOAD else 1.0
        for s, run in (("CT1", run1), ("CT2", run2)):
            out[s] = dict(op=np.where(run, u, 0.0), load=np.where(run, Ptot * share * ct_part, 0.0),
                          hi=np.where(run, hi * share, 0.0))
    else:
        s = W.stacks(key)[0]
        out[s] = dict(op=np.where(on, u, 0.0), load=np.where(on, Ptot, 0.0), hi=np.where(on, hi, 0.0))
    # substitute data: heat input (fuel flowmeter and data system outages) and NOx
    for s, d in out.items():
        sub = np.zeros(N, bool)
        fac = np.ones(N)
        for h0, n in r["hi_sub"][s]:
            sub[h0:h0 + n] = True
            fac[h0:h0 + n] = 1.06 + 0.08 * ((h0 * 7919 + n * 104729) % 1000) / 1000.0
        d["hi_sub"] = sub & (d["op"] > 0)
        d["hi_rep"] = np.where(d["hi_sub"], d["hi"] * fac, d["hi"])
        nsub = np.zeros(N, bool)
        for h0, n in r["nox_sub"][s]:
            nsub[h0:h0 + n] = True
        d["nox_sub"] = nsub & (d["op"] > 0)
    return out


def all_stacks(world):
    if "_stacks" not in world:
        world["_stacks"] = {k: stack_series(world, k) for k in W.KEYS}
    return world["_stacks"]


# ------------------------------------------------------------------------------------------------ EPA hourly
CAMD_HEADER = ["State", "Facility Name", "Facility ID", "Unit ID", "Date", "Hour", "Operating Time", "Gross Load (MW)",
               "SO2 Mass (lbs)", "SO2 Mass Measure Indicator", "NOx Rate (lbs/mmBtu)", "NOx Rate Measure Indicator",
               "NOx Mass (lbs)", "CO2 Mass (short tons)", "CO2 Mass Measure Indicator", "CO2 Rate (short tons/mmBtu)",
               "Heat Input (mmBtu)", "Heat Input Measure Indicator", "Primary Fuel Type", "Unit Type", "Program Code"]


def camd_rows(world):
    rng = np.random.default_rng(W.SEED + 101)
    stk = all_stacks(world)
    rows = []
    for key in sorted(W.KEYS, key=lambda k: (W.P[k]["state"], W.P[k]["epa"])):
        p = W.P[key]
        r = world["plants"][key]
        for s, d in stk[key].items():
            base_nox = 0.0072 + 0.0035 * rng.random()
            noxr = base_nox * rng.normal(1.0, 0.07, S.N)
            noxr = np.where(r["start"], rng.uniform(0.035, 0.07, S.N), noxr)
            noxr = np.where(d["nox_sub"], np.maximum(noxr, rng.uniform(0.012, 0.019)), noxr)
            for h in range(S.N):
                date = (W.FIRST_DAY + dt.timedelta(days=h // 24)).isoformat()
                uid = s
                if key in W.RENUMBER and date >= W.RENUMBER[key][2].isoformat():
                    uid = W.RENUMBER[key][1]
                op = d["op"][h]
                if op <= 0.0:
                    rows.append([p["state"], p["epa"], p["code"], uid, date, h % 24, "0.00"] + [""] * 11 +
                                ["Pipeline Natural Gas", "Combined cycle", "ARP"])
                    continue
                hi = d["hi_rep"][h]
                co2 = hi * W.CO2_PER_MMBTU
                rows.append([p["state"], p["epa"], p["code"], uid, date, h % 24, fmt(op, 2),
                             "%d" % int(round(d["load"][h])), fmt(hi * 0.0006, 2), "Calculated",
                             fmt(noxr[h], 4), "Substitute" if d["nox_sub"][h] else "Measured",
                             fmt(noxr[h] * hi, 1), fmt(co2, 1), "Calculated", "0.059", fmt(hi, 1),
                             "Substitute" if d["hi_sub"][h] else "Measured", "Pipeline Natural Gas", "Combined cycle",
                             "ARP"])
    return rows


def write_camd(world, path):
    return _w(path, CAMD_HEADER, camd_rows(world))


OVERRUN = ("MP", "OF", "KC")      # plants whose 2025 outage log carries a first estimate of return that ran long

PLAN_HISTORY = {("SB", "CT1"): " Fuel flowmeter replaced and recertified, March 2019.",
                ("SB", "CT2"): " Fuel flowmeter replaced and recertified, March 2019.",
                ("OF", "CT1"): " Initial certification of the NOx monitoring system, May 2021.",
                ("CB", "CT2"): " NOx analyzer replaced, October 2022.",
                ("LH", "CT1"): " Fuel flowmeter accuracy tests moved to the quarterly schedule, 2020."}


def write_plans(world, path):
    rows = []
    for key in sorted(W.KEYS, key=lambda k: W.P[k]["code"]):
        p = W.P[key]
        for s in W.stacks(key):
            note = "Gross load is the unit's combustion turbine output plus the steam turbine output apportioned to it."
            if p["config"] == "1x1ss":
                note = "Single-shaft combined cycle; gross load is the generator output."
            elif p["config"] == "1x1":
                note = "Gross load is the combustion turbine and steam turbine output."
            if key in W.TURBINE_ONLY_LOAD:
                note = "Gross load is the combustion turbine generator output only."
            db = (" Duct burner fuel is metered separately and included in the stack's heat input." if p["db"] else "")
            db += PLAN_HISTORY.get((key, s), "")
            if key in W.RENUMBER and s == W.RENUMBER[key][0]:
                old, new, when = W.RENUMBER[key]
                rows.append([p["code"], p["epa"], old, "CS0%s" % old[-1], "CC", "%d-%02d-01" % (p["year"], 6),
                             (when - dt.timedelta(days=1)).isoformat(), "HI: fuel flowmeter (Appendix D); CO2: Appendix G, calculated from heat input; NOx: CEMS",
                             note + db])
                rows.append([p["code"], p["epa"], new, "CS0%s" % old[-1], "CC", when.isoformat(), "",
                             "HI: fuel flowmeter (Appendix D); CO2: Appendix G, calculated from heat input; NOx: CEMS",
                             note + db + " Data acquisition and handling system replaced; unit ID changed from %s to %s "
                             "on %s." % (old, new, when.isoformat())])
                continue
            rows.append([p["code"], p["epa"], s, "CS0%s" % (s[-1] if s[-1].isdigit() else "1"), "CC",
                         "%d-%02d-01" % (p["year"], 6), "", "HI: fuel flowmeter (Appendix D); CO2: Appendix G, calculated from heat input; NOx: CEMS",
                         note + db])
    return _w(path, ["Facility ID", "Facility Name", "Unit ID", "Stack ID", "Unit Type", "Begin Date", "End Date",
                     "Methodologies", "Comment"], rows)


# ------------------------------------------------------------------------------------------------ EIA generators
def generators(key):
    """EIA generators of the offered combined cycle: (generator id, prime mover, nameplate MW, summer MW, winter MW)"""
    p = W.P[key]
    pm, db = p["pmax"], p["db"]
    if p["config"] == "1x1ss":
        return [("1", "CS", round(pm * 1.06, 1), round(W.cap_unfired(pm, 95) * 0.98), round(pm * 1.02))]
    if p["config"] == "1x1":
        ct, st = 0.66 * pm, 0.34 * pm + db
        return [("CT1", "CT", round(ct * 1.08, 1), round(0.66 * W.cap_unfired(pm, 95)), round(ct * 1.03)),
                ("ST1", "CA", round(st * 1.05, 1), round(0.34 * W.cap_unfired(pm, 95) + db), round(0.34 * pm * 1.03 + db))]
    ct, st = 0.33 * pm, 0.34 * pm + db
    return [("CT1", "CT", round(ct * 1.08, 1), round(0.33 * W.cap_unfired(pm, 95)), round(ct * 1.03)),
            ("CT2", "CT", round(ct * 1.08, 1), round(0.33 * W.cap_unfired(pm, 95)), round(ct * 1.03)),
            ("ST1", "CA", round(st * 1.05, 1), round(0.34 * W.cap_unfired(pm, 95) + db), round(0.34 * pm * 1.03 + db))]


UTILITY_ID = {k: 60000 + (W.P[k]["code"] * 7) % 9973 for k in W.KEYS}


def eia860_rows():
    rows = []
    for key in sorted(W.KEYS, key=lambda k: W.P[k]["code"]):
        p = W.P[key]
        for gid, pm, name, summ, wint in generators(key):
            dbf = ("Y" if p["db"] else "N") if pm in ("CA", "CS") else None
            byp = p["bypass"] if pm == "CT" else None
            minload = round(name * (0.42 if pm in ("CT", "CS") else 0.35))
            rows.append([UTILITY_ID[key], p["seller"], p["code"], p["eia"], p["state"], p["county"], gid,
                         "Natural Gas Fired Combined Cycle", pm, "CC1" if pm != "CS" else None, "S", dbf, byp, name,
                         0.85 if pm != "CA" else 0.9, summ, wint, minload, 6, p["year"], "NG", "OP"])
        for gid, mw, yr in W.PEAKERS.get(key, []):
            rows.append([UTILITY_ID[key], p["seller"], p["code"], p["eia"], p["state"], p["county"], gid,
                         "Natural Gas Fired Combustion Turbine", "GT", None, "S", None, None, mw, 0.85, round(mw * 0.93),
                         round(mw * 1.04), round(mw * 0.5), 5, yr, "NG", "OP"])
    return rows


EIA860_HEADER = ["Utility ID", "Utility Name", "Plant Code", "Plant Name", "State", "County", "Generator ID",
                 "Technology", "Prime Mover", "Unit Code", "Ownership", "Duct Burners",
                 "Can Bypass Heat Recovery Steam Generator", "Nameplate Capacity (MW)", "Nameplate Power Factor",
                 "Summer Capacity (MW)", "Winter Capacity (MW)", "Minimum Load (MW)", "Operating Month",
                 "Operating Year", "Energy Source 1", "Status"]


def write_eia860(path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Operable"
    ws.append(["Form EIA-860 Data - Schedule 3, 'Generator Data' (Operable Units Only), 2025 (extract: plants in the "
               "Calder PUD 2026 baseload request for offers)"])
    ws.append(EIA860_HEADER)
    rows = eia860_rows()
    for r in rows:
        ws.append(r)
    wb.save(path)
    return len(rows)


def write_xwalk(path):
    rows = []
    for key in sorted(W.KEYS, key=lambda k: W.P[k]["code"]):
        p = W.P[key]
        gens = generators(key)
        stacks = W.stacks(key)
        for s in stacks:
            for gid, pm, name, _, _ in gens:
                if p["config"] == "2x1" and pm == "CT" and gid[-1] != s[-1]:
                    continue
                rows.append([p["state"], p["epa"], p["code"], s, gid, round(name), "Pipeline Natural Gas", p["state"],
                             p["eia"], p["code"], gid, name, pm, "NG", "1-1" if pm in ("CT", "CS") else "1-many",
                             "No"])
    return _w(path, ["CAMD_STATE", "CAMD_FACILITY_NAME", "CAMD_PLANT_ID", "CAMD_UNIT_ID", "CAMD_GENERATOR_ID",
                     "CAMD_NAMEPLATE_CAPACITY", "CAMD_FUEL_TYPE", "EIA_STATE", "EIA_PLANT_NAME", "EIA_PLANT_ID",
                     "EIA_GENERATOR_ID", "EIA_NAMEPLATE_CAPACITY", "EIA_UNIT_TYPE", "EIA_FUEL_TYPE", "MATCH_TYPE_GEN",
                     "PLANT_ID_CHANGE"], rows)


# ------------------------------------------------------------------------------------------------ EIA-923 (fuel, generation)
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
          "November", "December"]


def monthly(arr, year):
    out = []
    for m in range(1, 13):
        a = (dt.date(year, m, 1) - W.FIRST_DAY).days * 24
        b = ((dt.date(year + (m == 12), m % 12 + 1, 1)) - W.FIRST_DAY).days * 24
        out.append(float(arr[a:b].sum()))
    return out


def eia923_rows(world, year):
    rng = np.random.default_rng(W.SEED + 200 + year)
    rows = []
    for key in sorted(W.KEYS, key=lambda k: W.P[k]["code"]):
        p = W.P[key]
        r = world["plants"][key]
        hi = r["hi"]
        dbfuel = np.where(r["on"], r["fire"] * p["dbhr"] / 1000.0 * r["op"], 0.0)
        net = r["net"]
        lines = []
        if p["config"] == "1x1ss":
            lines.append(("CS", monthly(hi, year), monthly(net, year)))
        else:
            ctshare = 0.66
            gen_ct = monthly(np.where(r["on"], (r["Ptot"] - r["fire"]) * r["op"] * ctshare, 0.0) - 0.0, year)
            gen_tot = monthly(net, year)
            gross = monthly(r["gross_mwh"], year)
            ct_net = [g_ct * (n / g if g > 0 else 0.0) if g > 0 else 0.5 * n for g_ct, n, g in zip(gen_ct, gen_tot, gross)]
            ca_net = [n - c for n, c in zip(gen_tot, ct_net)]
            fuel_ct = [a - b for a, b in zip(monthly(hi, year), monthly(dbfuel, year))]
            lines.append(("CT", fuel_ct, ct_net))
            lines.append(("CA", monthly(dbfuel, year), ca_net))
        for pm, fuel, gen in lines:
            hc = [round(1.032 + 0.012 * rng.random(), 3) for _ in range(12)]
            qty = [round(f / h) for f, h in zip(fuel, hc)]
            rows.append((p, pm, qty, hc, [round(f) for f in fuel], [round(f) for f in fuel], [round(g, 3) for g in gen]))
        for gid, mw, _ in W.PEAKERS.get(key, []):
            hrs = [int(rng.integers(0, 6)) if m in (1, 2, 3, 4, 5, 10, 11) else int(rng.integers(8, 70)) for m in range(1, 13)]
            gen = [round(h * mw * rng.uniform(0.8, 0.95), 3) for h in hrs]
            fuel = [round(g * rng.uniform(10.6, 11.4)) for g in gen]
            hc = [round(1.032 + 0.012 * rng.random(), 3) for _ in range(12)]
            rows.append((p, "GT", [round(f / h) for f, h in zip(fuel, hc)], hc, fuel, fuel, gen))
    return rows


def write_eia923(world, path):
    from openpyxl import Workbook
    wb = Workbook()
    first = True
    n = 0
    for year in (2025, 2024):
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Page 1 Generation and Fuel %d" % year
        ws.append(["EIA-923 Monthly Generation and Fuel Consumption Time Series. %d %s" % (
            year, "Final" if year == 2024 else "Final Release (published 2026-09-24)")])
        ws.append(["Source: U.S. Energy Information Administration, Form EIA-923 'Power Plant Operations Report'. "
                   "Extract: plants in the Calder PUD 2026 baseload request for offers."])
        ws.append(["Quantities of natural gas in thousand cubic feet (Mcf); heat content in MMBtu per Mcf; net generation "
                   "in megawatthours."])
        ws.append([])
        head = ["Plant Id", "Combined Heat And Power Plant", "Plant Name", "Operator Name", "Plant State"]
        if year == 2025:
            head += ["Balancing Authority Code"]
        head += ["NERC Region", "Reported Prime Mover", "Reported Fuel Type Code", "Physical Unit Label"]
        head += ["Quantity %s" % m for m in MONTHS] + ["MMBtuPer_Unit_%s" % m for m in MONTHS] + \
            ["Tot_MMBtu%s" % m for m in MONTHS] + ["Elec_MMBtu%s" % m for m in MONTHS] + \
            ["Netgen_%s" % m for m in MONTHS] + ["Total Fuel Consumption Quantity", "Total Fuel Consumption MMBtu",
                                                  "Elec Fuel Consumption MMBtu", "Net Generation (Megawatthours)", "YEAR"]
        ws.append(head)
        for p, pm, qty, hc, tot, elec, gen in eia923_rows(world, year):
            row = [p["code"], "N", p["eia"], p["seller"], p["state"]]
            if year == 2025:
                row += ["CPUD" if p["state"] == "WA" else "BPAT"]
            row += ["WECC", pm, "NG", "mcf"] + qty + hc + tot + elec + gen + [sum(qty), sum(tot), sum(elec),
                                                                               round(sum(gen), 3), year]
            ws.append(row)
            n += 1
    wb.save(path)
    return n


# ------------------------------------------------------------------------------------------------ offers
SUBSTATION = {"TL": "Tamarack Landing 500 kV", "JD": "Juniper Draw 230 kV", "SB": "Sorrel Butte 230 kV",
              "CB": "Coldbrook 230 kV", "MP": "Moraine Point 230 kV", "OF": "Ochre Flat 230 kV",
              "LH": "Lindgren Hill 230 kV", "AG": "Ashby Gulch 230 kV", "KC": "Kessling Creek 230 kV",
              "PM": "Brenmoor 115 kV"}


def write_offers(world, path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Offers received"
    ws.append(["Calder County Public Utility District No. 1 - 2026 Baseload Request for Offers - offers received "
               "(closed 30 September 2026, 5:00 p.m. Pacific)"])
    ws.append(["Offer ID", "Respondent", "Resource", "Resource description (respondent's)", "County", "State",
               "EIA plant code", "Contract Capacity (MW)", "Delivery point", "Term start", "Term end",
               "Capacity price ($/kW-month, 2027)", "Escalation (%/yr)", "Received"])
    for key in OFFER_ORDER:
        p = W.P[key]
        r = world["plants"][key]
        desc = {"2x1": "two-on-one combined cycle", "1x1": "one-on-one combined cycle",
                "1x1ss": "single-shaft combined cycle"}[p["config"]]
        if p["db"]:
            desc += " with supplementary firing"
        ws.append([OFFER_ID[key], p["seller"], p["trade"], desc, p["county"], p["state"],
                   None if key in W.OFFER_NO_CODE else p["code"], int(r["offer"]),
                   SUBSTATION[key],
                   "2027-01-01", "2036-12-31", p["price"], 2.0 if key not in ("PM", "CB") else 2.5, RECEIVED[key]])
    wb.save(path)
    return len(OFFER_ORDER)


# ------------------------------------------------------------------------------------------------ revenue meters
def local_starts():
    """for each hour of the record (standard time index), its start in local prevailing time with the UTC offset"""
    out = []
    t0 = dt.datetime.combine(W.FIRST_DAY, dt.time(), tzinfo=PST)
    for h in range(S.N):
        loc = (t0 + dt.timedelta(hours=h)).astimezone(LA)
        off = loc.utcoffset()
        sign = "-" if off < dt.timedelta(0) else "+"
        hh = abs(int(off.total_seconds() // 3600))
        out.append(loc.strftime("%Y-%m-%d %H:%M") + "%s%02d:00" % (sign, hh))
    return out


def write_meter(world, path):
    ls = local_starts()
    rows = []
    for key in OFFER_ORDER:
        net = world["plants"][key]["net"]
        res = W.P[key]["trade"]
        for h in range(S.N):
            rows.append([OFFER_ID[key], res, ls[h], 60, fmt(net[h], 3)])
    return _w(path, ["offer_id", "resource", "interval_start", "interval_minutes", "net_mwh"], rows)


# ------------------------------------------------------------------------------------------------ GADS events
CAUSE = {"plant": [("3611", "Main transformer"), ("9130", "Lightning"), ("3410", "Generator main leads"),
                   ("4000", "Heat recovery steam generator tube leak"), ("3899", "Other switchyard equipment")],
         "ct": [("5000", "Combustion turbine compressor"), ("5110", "Combustion turbine fuel system"),
                ("5250", "Combustion turbine controls"), ("5099", "Combustion turbine inlet air system")],
         "st": [("4300", "Steam turbine trip, high vibration"), ("4260", "Steam turbine lube oil system"),
                ("4400", "Condenser tube leak")],
         "SF": [("5260", "Combustion turbine failure to start, ignition"), ("5255", "Start-up control fault")]}


def gads_records(world, key):
    """the plant's 2025 GADS event lines, unit by unit, with the times the seller's system holds (UTC)"""
    rng = np.random.default_rng(W.SEED + 300 + W.P[key]["code"] % 997)
    p = W.P[key]
    r = world["plants"][key]
    units = W.gads_units(key)
    y0, y1 = dt.datetime(2025, 1, 1), dt.datetime(2026, 1, 1)
    lines = []

    def emit(unit, typ, s, e, avail, cause, desc, entered):
        lines.append(dict(unit=unit, type=typ, start=s, end=e, avail=avail, cause=cause, desc=desc, entered=entered))

    cap = {u: g[2] for u, g in zip(units, [(gg[0], gg[1], gg[2]) for gg in generators(key)])}
    if p["config"] == "1x1ss":
        cap = {"CS1": generators(key)[0][2]}
    else:
        gens = generators(key)
        cap = {}
        for gid, pm, name, summ, wint in gens:
            cap[gid] = summ
    revise = []
    for e in r["events"]:
        if not (e["end"] > y0 and e["start"] < y1):
            continue
        typ, scope = e["type"], e["scope"]
        if typ in S.FORCED:
            kind = "SF" if typ == "SF" else scope
            cc, cdesc = CAUSE[kind][int(rng.integers(0, len(CAUSE[kind])))]
        elif typ == "PO":
            cc, cdesc = "4999", "Annual planned outage, hot gas path and steam path inspections"
        else:
            cc, cdesc = "5080", "Maintenance outage, combustion turbine inlet and filter work"
        entered = e["start"] + dt.timedelta(minutes=int(rng.integers(20, 300)))
        affected = list(e["units"])
        induced = []
        if typ in S.FORCED and p["config"] != "1x1ss":
            if scope == "st" and p["bypass"] == "N":
                induced = [u for u in units if u.startswith("CT")]
            elif scope == "ct" and p["config"] == "1x1":
                induced = ["ST1"]
        for uu in affected + induced:
            emit(uu, typ, e["start"], e["end"], 0.0, cc, cdesc if uu in affected else "Unit out, forced by %s" % cdesc.lower(),
                 entered)
        if typ in S.FORCED and scope == "ct" and p["config"] == "2x1":
            # the steam turbine runs on the other turbine's steam: a forced derate
            emit("ST1", "D1", e["start"], e["end"], round(cap["ST1"] * 0.5, 1), cc, "Steam turbine derated, one heat "
                 "recovery steam generator out of service", entered)
        if typ in S.FORCED and (e["end"] - e["start"]).total_seconds() > 8 * 3600 and y0 <= e["start"] < y1:
            revise.append(e)
    # an outage first entered with a long estimate of its return to service, re-entered with its actual end once the
    # unit was back; the first entry stays in the seller's file (the estimate runs past the plant's next forced outage)
    over = None
    if key in OVERRUN:
        fe = sorted((e for e in r["events"] if e["type"] in S.FORCED and y0 <= e["start"] < y1), key=lambda e: e["start"])
        for a_, b_ in zip(fe, fe[1:]):
            if set(a_["units"]) & set(b_["units"]) and 2.0 <= (b_["start"] - a_["end"]).total_seconds() / 86400.0 <= 12.0:
                over = (a_, b_)
    # an outage extended past its first estimate is resubmitted with the new end under a new event number; the first
    # submission stays in the seller's file
    rng.shuffle(revise)
    nrev = max(1, int(round(len(revise) * 0.45))) if revise else 0
    for e in revise[:nrev]:
        dur = (e["end"] - e["start"]).total_seconds()
        est = e["start"] + dt.timedelta(seconds=int(dur * rng.uniform(0.3, 0.7) // 60 * 60))
        skip = over is not None and e is over[0]       # that outage's re-entry is the long estimate below
        first = [ln for ln in lines if ln["start"] == e["start"] and ln["end"] == e["end"] and ln["type"] == e["type"]]
        for ln in first:
            later = int(rng.integers(30, 180))
            if skip:
                continue
            emit(ln["unit"], ln["type"], ln["start"], est, ln["avail"], ln["cause"], ln["desc"],
                 ln["entered"] - dt.timedelta(minutes=5))
            ln["entered"] = est + dt.timedelta(minutes=later)
    if over is not None:
        a_, b_ = over
        est = (b_["start"] + dt.timedelta(days=1)).replace(hour=8, minute=0)
        first = [ln for ln in lines if ln["start"] == a_["start"] and ln["end"] == a_["end"] and ln["type"] == a_["type"]]
        for ln in first:
            emit(ln["unit"], ln["type"], ln["start"], est, ln["avail"], ln["cause"], ln["desc"], ln["entered"])
            ln["entered"] = a_["end"] + dt.timedelta(minutes=133)
    # reserve shutdowns from the dispatch: each economic shutdown in 2025, on every unit
    on = r["on"]
    op = r["op"]
    t0 = dt.datetime.combine(W.FIRST_DAY, dt.time())
    outage_hours = r["full_out"] > 0.0
    h = W.H2025
    while h < S.N:
        if not on[h] and not outage_hours[h]:
            a = h
            while h < S.N and not on[h]:
                h += 1
            prev = a - 1
            s = t0 + dt.timedelta(hours=prev) + dt.timedelta(minutes=int(round(op[prev] * 60))) if prev >= 0 and on[prev] \
                else t0 + dt.timedelta(hours=a)
            if h < S.N:
                e = t0 + dt.timedelta(hours=h) + dt.timedelta(minutes=int(round((1.0 - op[h]) * 60)))
            else:
                e = t0 + dt.timedelta(hours=h)
            if any(outage_hours[a:h]):
                continue
            lag = int(rng.integers(3, 140))
            for uu in units:
                emit(uu, "RS", s, e, 0.0, "0000", "Reserve shutdown", s + dt.timedelta(minutes=lag))
        else:
            h += 1
    # event numbers, unit by unit, in order of entry; the log lists each unit's events by number
    num = {}
    for ln in sorted(lines, key=lambda ln: (ln["unit"], ln["entered"], ln["start"])):
        num[ln["unit"]] = num.get(ln["unit"], 0) + 1
        ln["event"] = num[ln["unit"]]
    lines.sort(key=lambda ln: (ln["unit"], ln["event"]))
    return lines


def to_utc(t):
    return (t.replace(tzinfo=PST)).astimezone(UTC)


def write_gads(world, path):
    rows = []
    for key in OFFER_ORDER:
        p = W.P[key]
        for ln in gads_records(world, key):
            rows.append([OFFER_ID[key], p["trade"], ln["unit"], "%03d" % ln["event"], ln["type"],
                         to_utc(ln["start"]).strftime("%Y-%m-%d %H:%M"), to_utc(ln["end"]).strftime("%Y-%m-%d %H:%M"),
                         fmt(ln["avail"], 1), ln["cause"], ln["desc"], to_utc(ln["entered"]).strftime("%Y-%m-%d %H:%M")])
    return _w(path, ["offer_id", "resource", "unit", "event_number", "event_type", "start_utc", "end_utc",
                     "net_available_capacity_mw", "cause_code", "description", "entered_utc"], rows)


# ------------------------------------------------------------------------------------------------ 2027 planned outages
def write_outages27(world, path):
    rng = np.random.default_rng(W.SEED + 500)
    rows = []
    for key in OFFER_ORDER:
        p = W.P[key]
        units = [u for u in W.gads_units(key)]
        s = dt.date(2027, 4, 1) + dt.timedelta(days=int(rng.integers(0, 45)))
        d = int(rng.integers(10, 24))
        work = "major inspection" if rng.random() < 0.3 else "hot gas path inspection"
        rows.append([OFFER_ID[key], p["trade"], ";".join(units), s.isoformat(), (s + dt.timedelta(days=d)).isoformat(),
                     "PO", "%s, %s" % (units[0], work)])
        if rng.random() < 0.5:
            s2 = dt.date(2027, 10, 1) + dt.timedelta(days=int(rng.integers(0, 40)))
            rows.append([OFFER_ID[key], p["trade"], ";".join(units), s2.isoformat(),
                         (s2 + dt.timedelta(days=int(rng.integers(2, 6)))).isoformat(), "MO", "Inlet filter replacement"])
    return _w(path, ["offer_id", "resource", "units", "planned_start", "planned_end", "event_type", "work"], rows)
