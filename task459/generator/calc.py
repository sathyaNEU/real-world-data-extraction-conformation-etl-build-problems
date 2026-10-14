"""Readers of the shipped pack and every route the design names, from the files alone.

The golden: each offered plant's pounds of CO2 per net MWh delivered at its contract capacity, read from the hours
(2024 and 2025) in which its revenue meter shows it delivering that capacity, against EPA's hourly CO2 for the same
hours; a plant complies at 925 or below. The other routes: the 2025 annual average (the default), the same without
the start-up and shutdown hours, and the plant's highest-output hours."""
import csv
import datetime as dt
import os
import re
from zoneinfo import ZoneInfo

import numpy as np

LA = ZoneInfo("America/Los_Angeles")
PST = dt.timezone(dt.timedelta(hours=-8))
FIRST = dt.date(2024, 1, 1)
N = ((dt.date(2025, 12, 31) - FIRST).days + 1) * 24
H2025 = (dt.date(2025, 1, 1) - FIRST).days * 24
STANDARD = 925.0
FORCED = ("U1", "U2", "U3", "SF")


def find(target, pattern):
    for f in sorted(os.listdir(target)):
        if re.search(pattern, f):
            return os.path.join(target, f)
    raise FileNotFoundError(pattern)


def read_csv(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


# ---------------------------------------------------------------------------------------------- offers and plants
def offers(target):
    from openpyxl import load_workbook
    wb = load_workbook(find(target, r"^rfo2026_offers_received\.xlsx$"), read_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    head = list(rows[1])
    out = []
    for r in rows[2:]:
        if not r or r[0] is None:
            continue
        d = dict(zip(head, r))
        out.append(d)
    return out


def plant_names(target):
    """EIA plant code -> (EIA plant name, county, state) from the generator inventory"""
    from openpyxl import load_workbook
    wb = load_workbook(find(target, r"^eia860_.*\.xlsx$"), read_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    head = list(rows[1])
    out = {}
    for r in rows[2:]:
        d = dict(zip(head, r))
        out[int(d["Plant Code"])] = (d["Plant Name"], d["County"], d["State"])
    return out


def offer_plant_codes(target):
    """offer id -> EIA plant code: the offer's code, or for an offer that left it blank the plant in the generator
    inventory whose name shares the offer's resource name's first two words, in the offer's county"""
    names = plant_names(target)
    out = {}
    for o in offers(target):
        code = o["EIA plant code"]
        if code in (None, ""):
            stem = " ".join(str(o["Resource"]).split()[:2]).lower()
            hits = [c for c, (nm, county, st) in names.items()
                    if nm.lower().startswith(stem) and county == o["County"] and st == o["State"]]
            assert len(hits) == 1, (o["Offer ID"], hits)
            code = hits[0]
        out[o["Offer ID"]] = int(code)
    return out


# ---------------------------------------------------------------------------------------------- hourly records
def std_index(date, hour):
    return (dt.date.fromisoformat(date) - FIRST).days * 24 + int(hour)


def camd(target):
    """per facility (EIA plant code): CO2 (short tons), operating (any stack), every running stack at full operating
    time, heat-input substitute and NOx substitute flags (any operating stack), per standard-time hour"""
    out = {}
    with open(find(target, r"^camd_hourly_emissions"), newline="", encoding="utf-8") as f:
        rd = csv.reader(f)
        head = next(rd)
        ix = {h: i for i, h in enumerate(head)}
        for r in rd:
            fac = int(r[ix["Facility ID"]])
            d = out.get(fac)
            if d is None:
                d = out[fac] = dict(co2=np.zeros(N), op=np.zeros(N, bool), partial=np.zeros(N, bool),
                                    hisub=np.zeros(N, bool), noxsub=np.zeros(N, bool), stackhisub=np.zeros(N, int),
                                    units=set())
            h = std_index(r[ix["Date"]], r[ix["Hour"]])
            opt = float(r[ix["Operating Time"]])
            d["units"].add(r[ix["Unit ID"]])
            if opt <= 0:
                continue
            d["op"][h] = True
            if opt < 1.0:
                d["partial"][h] = True
            d["co2"][h] += float(r[ix["CO2 Mass (short tons)"]])
            if r[ix["Heat Input Measure Indicator"]] == "Substitute":
                d["hisub"][h] = True
                d["stackhisub"][h] += 1
            if r[ix["NOx Rate Measure Indicator"]] == "Substitute":
                d["noxsub"][h] = True
            d.setdefault("co2_ind", set()).add(r[ix["CO2 Mass Measure Indicator"]])
    return out


T0 = dt.datetime(2024, 1, 1, 8, tzinfo=dt.timezone.utc)        # 1 January 2024, midnight Pacific standard time


def meter_index(stamp, ignore_offset=False):
    """the standard-time hour index of an interval start written as local time with its UTC offset"""
    t = dt.datetime.fromisoformat(stamp)
    if ignore_offset:
        t = t.replace(tzinfo=PST)
    return int(round((t - T0).total_seconds() / 3600.0))


def meter(target, ignore_offset=False):
    """offer id -> net MWh per standard-time hour, from the interval starts in local time with their offsets"""
    rows = read_csv(find(target, r"^rfo2026_offer_meter_data"))
    out = {}
    for r in rows:
        h = meter_index(r["interval_start"], ignore_offset)
        arr = out.get(r["offer_id"])
        if arr is None:
            arr = out[r["offer_id"]] = np.full(N, np.nan)
        if 0 <= h < N:
            arr[h] = float(r["net_mwh"])
    return out


def load(target):
    o = offers(target)
    codes = offer_plant_codes(target)
    return dict(offers=o, codes=codes, camd=camd(target), meter=meter(target),
                cap={x["Offer ID"]: float(x["Contract Capacity (MW)"]) for x in o},
                price={x["Offer ID"]: float(x["Capacity price ($/kW-month, 2027)"]) for x in o},
                name={x["Offer ID"]: x["Resource"] for x in o})


NEED = 1000.0


def slate(L, comply, need=NEED):
    """the need filled from the complying offers in order of capacity price, lowest first, each offer taken whole and
    any offer passed over that would take the total above the need: (taken, total MW, passed over in price order)"""
    total, taken, passed = 0.0, [], []
    for o in sorted(comply, key=lambda o: (L["price"][o], o)):
        if total + L["cap"][o] <= need + 1e-9:
            taken.append(o)
            total += L["cap"][o]
        else:
            passed.append(o)
    return taken, total, passed


# ---------------------------------------------------------------------------------------------- routes
def rate(co2, net, sel):
    return float(co2[sel].sum() * 2000.0 / net[sel].sum())


def series(L, oid):
    c = L["camd"][L["codes"][oid]]
    return c, L["meter"][oid]


def at_capacity(L, oid, band=0.02, years=(2024, 2025)):
    """the golden: hours whose metered net output is within band of the contract capacity, every running stack at
    full operating time"""
    c, net = series(L, oid)
    cap = L["cap"][oid]
    idx = np.arange(N)
    yr = np.zeros(N, bool)
    if 2024 in years:
        yr |= idx < H2025
    if 2025 in years:
        yr |= idx >= H2025
    sel = yr & c["op"] & ~c["partial"] & (np.abs(net - cap) <= band * cap)
    return rate(c["co2"], net, sel), int(sel.sum())


def fit_at_capacity(L, oid, degree=2, lo=0.8):
    """a curve: hourly CO2 on hourly net output over full operating hours from lo of the contract capacity upward,
    evaluated at the contract capacity"""
    c, net = series(L, oid)
    cap = L["cap"][oid]
    sel = c["op"] & ~c["partial"] & (net >= lo * cap) & (net <= 1.06 * cap)
    co = np.polyfit(net[sel] / cap, c["co2"][sel] * 2000.0, degree)
    return float(np.polyval(co, 1.0) / cap)


def annual(L, oid, year=2025):
    c, net = series(L, oid)
    idx = np.arange(N)
    sel = (idx >= H2025) if year == 2025 else ((idx < H2025) if year == 2024 else np.ones(N, bool))
    return rate(c["co2"], net, sel)


def annual_no_starts(L, oid):
    c, net = series(L, oid)
    idx = np.arange(N)
    sel = (idx >= H2025) & c["op"] & ~c["partial"]
    prev = np.concatenate([[False], c["op"][:-1]])
    nxt = np.concatenate([c["op"][1:], [False]])
    sel &= prev & nxt
    return rate(c["co2"], net, sel)


def top_output(L, oid, q=99, frac=0.95):
    c, net = series(L, oid)
    full = c["op"] & ~c["partial"]
    top = np.percentile(net[full], q)
    sel = full & (net >= frac * top)
    return rate(c["co2"], net, sel)


GOLDEN = ("band2", "band1", "band3", "band5", "band2_2025", "fit2", "fit1")


def low_output_excess(L, oid, frac=0.75, min_share=0.2):
    """a plant's rate in its full operating hours below frac of its Contract Capacity against its rate at it, per
    cent above; None where such hours are under min_share of its full operating hours (the base-loaded plants)"""
    c, net = series(L, oid)
    cap = L["cap"][oid]
    full = c["op"] & ~c["partial"] & (net > 0)
    lo = full & (net < frac * cap)
    if lo.sum() < min_share * full.sum():
        return None
    at = full & (np.abs(net - cap) <= 0.02 * cap)
    return 100.0 * (rate(c["co2"], net, lo) / rate(c["co2"], net, at) - 1.0)


def golden_family(L, oid):
    out = {}
    out["band2"] = at_capacity(L, oid, 0.02)[0]
    out["band1"] = at_capacity(L, oid, 0.01)[0]
    out["band3"] = at_capacity(L, oid, 0.03)[0]
    out["band5"] = at_capacity(L, oid, 0.05)[0]
    out["band2_2025"] = at_capacity(L, oid, 0.02, years=(2025,))[0]
    out["fit2"] = fit_at_capacity(L, oid, 2)
    out["fit1"] = fit_at_capacity(L, oid, 1, lo=0.9)
    return out


def routes(L):
    R = {}
    for oid in sorted(L["cap"]):
        R[oid] = dict(golden=at_capacity(L, oid)[0], n=at_capacity(L, oid)[1], default=annual(L, oid, 2025),
                      both_years=annual(L, oid, None), no_starts=annual_no_starts(L, oid), top=top_output(L, oid),
                      family=golden_family(L, oid))
    return R


def signs(R, key):
    return sorted(o for o, r in R.items() if (r[key] if key in r else r["family"][key]) <= STANDARD)


# ---------------------------------------------------------------------------------------------- Ask A: forced outages
def gads(target):
    return read_csv(find(target, r"^rfo2026_offer_gads_events"))


def _merge(iv):
    iv = sorted(iv)
    out = []
    for a, b in iv:
        if out and a <= out[-1][1]:
            out[-1][1] = max(out[-1][1], b)
        else:
            out.append([a, b])
    return out


def _intersect(a, b):
    out, i, j = [], 0, 0
    while i < len(a) and j < len(b):
        lo, hi = max(a[i][0], b[j][0]), min(a[i][1], b[j][1])
        if lo < hi:
            out.append([lo, hi])
        if a[i][1] < b[j][1]:
            i += 1
        else:
            j += 1
    return out


Y0 = dt.datetime(2025, 1, 1, tzinfo=PST)
Y1 = dt.datetime(2026, 1, 1, tzinfo=PST)


def utc(s):
    return dt.datetime.strptime(s, "%Y-%m-%d %H:%M").replace(tzinfo=dt.timezone.utc)


def current_entries(rows):
    """a unit's entries for one outage share its start; the latest entered stands"""
    best = {}
    for r in rows:
        k = (r["offer_id"], r["unit"], r["start_utc"])
        if k not in best or r["entered_utc"] > best[k]["entered_utc"]:
            best[k] = r
    return list(best.values())


def forced_outages(target):
    """offer id -> (forced outage hours in 2025, forced outages in 2025): the whole plant out on a forced event when
    every combustion turbine (or the single-shaft unit) is out on one, each unit's current entries merged in time"""
    allrows = gads(target)
    rows = current_entries(allrows)
    by = {}
    for r in rows:
        if r["event_type"] not in FORCED:
            continue
        by.setdefault(r["offer_id"], {}).setdefault(r["unit"], []).append((utc(r["start_utc"]), utc(r["end_utc"])))
    units = {}
    for r in allrows:
        units.setdefault(r["offer_id"], set()).add(r["unit"])
    out = {}
    for oid in sorted(units):
        drivers = sorted(u for u in units[oid] if u.startswith("CT") or u.startswith("CS"))
        iv = None
        for u in drivers:
            m = _merge(by.get(oid, {}).get(u, []))
            iv = m if iv is None else _intersect(iv, m)
        iv = _merge(iv or [])
        hours, count = 0.0, 0
        for a, b in iv:
            lo, hi = max(a, Y0), min(b, Y1)
            if lo < hi:
                hours += (hi - lo).total_seconds() / 3600.0
                count += 1
        out[oid] = (hours, count)
    return out


def forced_outages_naive(target, mode="sum"):
    """the reflex routes, on the dictionary's Pacific year: 'sum' every forced line's duration within 2025, per offer,
    counting lines; 'union' the union of every unit's forced records, counting the merged spans"""
    rows = gads(target)
    out = {}
    if mode == "sum":
        for r in rows:
            if r["event_type"] not in FORCED:
                continue
            a, b = utc(r["start_utc"]), utc(r["end_utc"])
            lo, hi = max(a, Y0), min(b, Y1)
            h, c = out.get(r["offer_id"], (0.0, 0))
            if lo < hi:
                out[r["offer_id"]] = (h + (hi - lo).total_seconds() / 3600.0, c + 1)
        return out
    by = {}
    for r in rows:
        if r["event_type"] in FORCED:
            by.setdefault(r["offer_id"], []).append((utc(r["start_utc"]), utc(r["end_utc"])))
    for oid, iv in by.items():
        hours, count = 0.0, 0
        for a, b in _merge(iv):
            lo, hi = max(a, Y0), min(b, Y1)
            if lo < hi:
                hours += (hi - lo).total_seconds() / 3600.0
                count += 1
        out[oid] = (hours, count)
    return out


# ---------------------------------------------------------------------------------------------- Ask B: substitute data
def substitute_hours(L):
    """offer id -> 2025 hours in which any operating stack's heat input, and with it its CO2, came from substitute
    data"""
    out = {}
    for oid, code in L["codes"].items():
        c = L["camd"][code]
        sel = np.arange(N) >= H2025
        out[oid] = int((sel & c["op"] & c["hisub"]).sum())
    return out


def substitute_naive(L):
    out = {}
    for oid, code in L["codes"].items():
        c = L["camd"][code]
        sel = np.arange(N) >= H2025
        out[oid] = dict(stack_hours=int(c["stackhisub"][sel].sum()), any_flag=int((sel & c["op"] & (c["hisub"] | c["noxsub"])).sum()),
                        co2_indicator=0 if c.get("co2_ind") == {"Calculated"} else -1)
    return out
