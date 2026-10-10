"""Constructed documents of the pack: the vulnerability standard, the SRE maintenance standard, the
provider service schedule and window calendar, the Q3 close-out, the planning thread, the data
dictionary, the provenance record and the superseded CVSS-band memo (a declared distractor).

All documents are short and functional. Container metadata (producer signature, build-clock dates)
is scrubbed in build.py after they are written.
"""
import datetime as dt

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
import docx
from docx.shared import Pt

from params import ORG, PROVIDER, PEOPLE, LABEL, CLOUD, COLO, AS_OF, TARGET_DAYS, THRESH, RACK
from cloud import MONTHS_Q3

STANDARD = "vulnerability_management_standard.docx"
SRE = "sre_maintenance_standard.pdf"
SCHEDULE = "colocation_service_schedule_and_window_calendar.pdf"
CLOSEOUT = "q3_2026_remediation_closeout.pdf"
THREAD = "november_round_planning_thread.eml"
DICT = "warehouse_data_dictionary.md"
PROVENANCE = "extract_provenance_2026-10-23.md"
MEMO = "superseded_cvss_band_allocation_memo.pdf"
MONTHNAME = {7: "July", 8: "August", 9: "September"}


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=10, leading=13.5,
                          spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=14,
                        leading=18, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=10.5,
                        leading=14, spaceBefore=7, spaceAfter=3)
    small = ParagraphStyle("s", parent=body, fontSize=8.5, leading=11,
                           textColor=colors.HexColor("#444444"))
    return body, h1, h2, small


def _pdf(path, title, story, footer):
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
                            topMargin=18 * mm, bottomMargin=20 * mm, title=title, author=ORG,
                            subject="", creator=ORG, invariant=1)

    def on_page(c, d):
        c.saveState()
        c.setFont("Helvetica", 7.5)
        c.setFillColor(colors.HexColor("#555555"))
        c.drawString(20 * mm, 12 * mm, footer)
        c.drawRightString(190 * mm, 12 * mm, f"Page {d.page}")
        c.restoreState()
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)


def write_sre(path):
    body, h1, h2, small = _styles()
    P = lambda t, s=body: Paragraph(t, s)
    st = [P(ORG + " / Site Reliability Engineering", small),
          P("Production maintenance standard", h1),
          P("Owner: SRE lead. Applies to all production estates.", small), Spacer(1, 6),
          P("1 Failure-domain headroom", h2),
          P("Maintenance must never take an estate below one failure domain of headroom at its "
            "forecast peak. A failure domain is one rack."),
          P("2 Colocated estates", h2),
          P("Work on a colocated estate happens only by draining hosts through the provider. While "
            "an estate is being drained it keeps at least one rack of hosts in hand above the "
            "transactions its forecast peak needs, over the hours the window runs."),
          P("3 Cloud estates", h2),
          P("Cloud estates are patched in place by the crews, package by package. No draining is "
            "involved.")]
    _pdf(path, "SRE maintenance standard", st, f"{ORG}  |  SRE  |  Maintenance standard")


def write_schedule(path, W):
    body, h1, h2, small = _styles()
    P = lambda t, s=body: Paragraph(t, s)
    from params import NOV_OFFICE, WINDOWS
    st = [P(PROVIDER, small), P("Colocation service schedule and November window calendar", h1),
          P("Issued to " + ORG + " by " + PEOPLE["provider"] + ", service delivery manager.", small),
          Spacer(1, 6),
          P("1 Maintenance windows", h2),
          P("Payments windows run Tuesday and Thursday, 20:00 to 24:00 local. Checkout windows run "
            "Monday and Wednesday, 22:00 to 02:00 local; a checkout window crosses local midnight "
            "and spans two dates."),
          P("2 Draining", h2),
          P("Within a window hosts are drained one cycle at a time, a cycle every 30 minutes, so a "
            "four-hour window runs eight cycles. The number of hosts drained at once is agreed per "
            "window. Each drained host is taken out of the serving pool, serviced and returned to "
            "the pool."),
          P("3 Requests", h2),
          P("A team requests a set of hosts for a window. The provider acknowledges each request "
            "with the hosts accepted. Requests are taken in the order they are submitted.")]
    cal = [["Estate", "Window dates (November)"]]
    cal.append(["Payments", ", ".join(d.strftime("%a %d %b") for d in NOV_OFFICE["payments"])])
    cal.append(["Checkout", ", ".join(d.strftime("%a %d %b") for d in NOV_OFFICE["checkout"])])
    t = Table(cal, colWidths=[35 * mm, 125 * mm])
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 9),
                           ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 9),
                           ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.HexColor("#888888")),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
    st += [P("4 November windows booked to " + ORG, h2), t]
    _pdf(path, "Colocation service schedule", st,
         f"{PROVIDER}  |  Service schedule  |  November 2026")


def write_closeout(path, W):
    body, h1, h2, small = _styles()
    P = lambda t, s=body: Paragraph(t, s)
    truth = W["co_truth"]
    st = [P(ORG + " / Vulnerability management office", small),
          P("Q3 2026 remediation close-out", h1),
          P("Cloud estates only. Prepared by " + PEOPLE["planner"] + ", 7 October 2026.", small),
          Spacer(1, 6),
          P("This close-out reports exploitable host exposures the crews' tickets took off the four "
            "cloud estates in the third quarter. An exposure is one exploitable vulnerability on one "
            "host. The colocated estates are closed by the provider and are not in this report.")]
    head = ["Estate"] + [MONTHNAME[m] for m in (7, 8, 9)] + ["Quarter"]
    rows = [head]
    gtot = 0
    for est in CLOUD:
        r = [LABEL[est]]
        s = 0
        for m in (7, 8, 9):
            v = truth.get((est, m), 0)
            r.append(f"{v:,}")
            s += v
        r.append(f"{s:,}")
        gtot += s
        rows.append(r)
    rows.append(["Total", "", "", "", f"{gtot:,}"])
    t = Table(rows, colWidths=[40 * mm] + [30 * mm] * 4)
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 9),
                           ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 9),
                           ("FONT", (0, -1), (-1, -1), "Helvetica-Bold", 9),
                           ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.HexColor("#888888")),
                           ("LINEABOVE", (0, -1), (-1, -1), 0.5, colors.HexColor("#888888")),
                           ("ALIGN", (1, 0), (-1, -1), "RIGHT"), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
    st += [Spacer(1, 4), t, Spacer(1, 6),
           P("Figures are exposures closed, counted on the exploit score each ticket carried on the "
             "day it was cut. This is the number the audit committee sees.", small)]
    _pdf(path, "Q3 2026 remediation close-out", st,
         f"{ORG}  |  Vulnerability management  |  Q3 2026 close-out")


def write_memo(path):
    body, h1, h2, small = _styles()
    P = lambda t, s=body: Paragraph(t, s)
    st = [P(ORG + " / Vulnerability management office", small),
          P("Patch allocation by CVSS band (superseded)", h1),
          P("Superseded by the November 2026 exposure-based allocation. Retained for reference "
            "only.", small), Spacer(1, 6),
          P("Until October 2026 the crews' monthly change tickets were spread evenly across the "
            "estates by CVSS band: a fixed share to each estate in proportion to its count of "
            "critical and high findings. From the November round this basis is replaced: tickets go "
            "where they take out the most exploitable host exposure in the month. The band split "
            "below no longer governs the allocation and is kept only to show the prior method."),
          P("Prior fixed shares: payments 18 per cent, checkout 17 per cent, search 16 per cent, "
            "media 20 per cent, internal tools 12 per cent, data pipeline 17 per cent.")]
    _pdf(path, "Superseded CVSS-band allocation memo", st,
         f"{ORG}  |  Superseded  |  CVSS-band allocation")


def write_standard(path):
    d = docx.Document()
    d.add_heading(ORG + " vulnerability management standard", level=0)
    d.add_paragraph("Version 3. Owner: vulnerability management office. In force from the November "
                    "2026 round.")
    d.add_heading("1 How tickets are allocated", level=1)
    d.add_paragraph("The crews raise 300 change tickets a month. From November 2026 tickets go "
                    "where they take out the most exploitable host exposure in the month. An "
                    "exposure is one exploitable vulnerability on one host; it is taken out when it "
                    "is gone from the host by month end.")
    d.add_heading("2 What counts as exploitable", level=1)
    d.add_paragraph("A finding is exploitable when its exploit-prediction score is at least "
                    f"{THRESH:.2f} on the day the ticket is cut. For the forward month the latest "
                    "score is used. The scanner's exploit-available flag is a convenience marker and "
                    "is not the measure.")
    d.add_heading("3 One ticket, one package", level=1)
    d.add_paragraph("A ticket updates one package. It is raised against the package whose "
                    "highest-scoring open finding on the hosts it names scores highest, and it names "
                    "the hosts that carry that package below its fixed version.")
    d.add_heading("4 Cloud and colocated estates", level=1)
    d.add_paragraph("A cloud ticket cut on the month's first Monday is closed in place by the crews "
                    "within the month. Colocated work is carried out by the provider under the "
                    "service schedule.")
    d.add_heading("5 Remediation target", level=1)
    d.add_paragraph(f"A ticket should run on its last host within {TARGET_DAYS} days of the vendor's "
                    "first release of the fix. A ticket counts as completed on the day its last "
                    "host ran it successfully.")
    d.add_heading("6 Scan coverage", level=1)
    d.add_paragraph("A cloud host is in service when its power state is running; standby and stopped "
                    "instances are not in service. A host is covered when it has an authenticated scan "
                    "dated on or after the export date less 14 days.")
    for p in d.paragraphs:
        for r in p.runs:
            r.font.size = r.font.size or Pt(10.5)
    d.save(path)


THREAD_TEXT = """From: {planner} <reina.guzman@sendalia.example>
To: Santiago Vidal <santiago.vidal@sendalia.example>; Prudencio Cañete <prudencio.canete@sendalia.example>; Fabio Montalbán <fabio.montalban@sendalia.example>; Eva Mas <eva.mas@sendalia.example>
Subject: November patch round - split by exposure
Date: Wed, 21 Oct 2026 10:14:00 +0200
Message-ID: <nov-round-1@sendalia.example>

All,

November is the first month we split the 300 tickets on exploitable exposure rather than spreading them by band. I sign the split on Friday 30 October and the crews key off it on Monday 2 November. I will bring one rate and the list, with the figures for Santiago.

Reina

--- Santiago Vidal, 21 Oct 2026 11:02 +0200 ---
Put the bulk on payments. That is where our worst exposure per host sits and the board expects to see it covered first.

--- Prudencio Cañete, 21 Oct 2026 12:20 +0200 ---
Just remember payments and checkout are colocated. We only touch those by draining hosts in the provider's windows, and we hold a rack above the forecast peak while we do. There is only so much we can drain in a month. The cloud estates we patch ourselves, package by package.

--- Felicia Infante (Centro de Datos Guadalhorce), 21 Oct 2026 14:05 +0200 ---
November windows are booked to you. The calendar and the service schedule are in the folder. Drained hosts come back serviced and in the pool, as always.

--- Fabio Montalbán, 21 Oct 2026 15:40 +0200 ---
The scanner export, the managed-host feed, the inventory, the capacity register, the forecast and the acknowledgements are all in the folder as at 23 October. The Q3 close-out covers the cloud estates only.

--- Eva Mas, 21 Oct 2026 16:12 +0200 ---
Crews are ready for whatever the split says on the 2nd.
""".format(planner=PEOPLE["planner"])


def write_thread(path):
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(THREAD_TEXT)


def write_dict(path):
    t = """# Warehouse data dictionary

Field meanings for the November round extracts. As at 23 October 2026.

## vuln_findings_2026-10-23.csv (office scanner, cloud estates)
- host_id: scanner host identifier.
- estate: one of search, media, tools, pipeline.
- package: the installed package the finding is against.
- cve: the vulnerability identifier. Join to epss_score_history_2026.parquet for the score.
- cvss_base: CVSS base score.
- exploit_available: the scanner's exploit-available marker. A convenience flag, not the office measure.
- first_seen, last_authenticated_scan: finding first observed, host last authenticated scan (date).

## epss_score_history_2026.parquet
- cve, score_date, epss: the daily exploit-prediction score for each cve. The November reading uses the latest score; a past ticket uses the score on its cut day.

## colocation_managed_host_findings_2026-10-23.jsonl (provider feed, colocated estates)
- host_id, package, cve, epss, first_seen: open findings on the provider-managed hosts, current state.

## colocation_host_inventory_2026-10-23.xlsx
- host_id, estate, pool, role: the colocated host and its role.
- in_service_since: the date the host last entered the serving pool.

## colocation_capacity_register_2026-10.xlsx
- hosts_in_service: hosts serving the estate that month.
- per_host_tps: transactions per second one host serves.
- failure_domain_hosts: hosts in one rack (one failure domain).
- drain_cycles_per_window: drain cycles in one window.

## november_window_calendar.csv
- estate, window_date, window_hours_local, drain_cycles: each November window booked to the office; window_hours_local is the local start and end, and a window ending after midnight runs into the next date.

## estate_transaction_forecast_2026.csv
- estate, date, hour, tps: forecast transactions per second, by hour.

## colocation_change_acknowledgements_may_oct_2026.xlsx
- request_id, estate, window_date, window_hours_local, team, submission_order, hosts_requested, hosts_accepted, status, reason: the provider's record of each change request and what it accepted.
- host_ids: the hosts requested. accepted_host_ids: the hosts the provider accepted and drained in that window.

## cloud_patch_ticket_log_q3_2026.csv / cloud_q3_closed_findings.csv
- the Q3 cloud tickets and the findings each closed. The close-out figures recompute from the fixed findings against the score on each ticket's cut date.

## crew_deployment_log_2026.csv
- ticket_id, estate, package, vendor_first_release, run_date, outcome, is_last_host: each host run of a deployment. A fix lands when a host runs it with outcome succeeded; a rolled_back run did not land.

## vendor_advisory_feed.json
- advisory, package, first_published, latest_revision: one record per vendor advisory; first_published is the vendor's first release of the fix, latest_revision the date it was last republished.

## cloud_asset_register_2026-10-23.csv
- instance_id: the stable host key. hostname: the current name. power_state: running, standby or stopped.

## scanner_coverage_2026-10.csv
- hostname, instance_id, estate, last_authenticated_scan: coverage keyed by the name in effect at the scan. Join on instance_id, which is stable across a rename.
"""
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(t)


def write_provenance(path, info):
    t = f"""# Extract provenance, November 2026 patch round

Prepared by {PEOPLE['data']} (data platform lead), 23 October 2026.

## About this folder

This folder is a constructed case. {ORG}, {PROVIDER}, every estate, host, person and figure in it were built for this case. CVE identifiers follow the public CVE format and the exploit-prediction scores follow the shape of FIRST's public EPSS feed; neither identifies a real vulnerability or a real score. All files may be used and shared under CC BY 4.0.

## Files

| File | What it is | How it was produced |
|---|---|---|
| vuln_findings_2026-10-23.csv | Office scanner open findings, cloud estates ({info['spine']:,} rows) | Scanner reporting export, 23 October 2026 |
| epss_score_history_2026.parquet | Daily exploit-prediction scores ({info['history']:,} rows) | Scoring feed snapshot |
| colocation_managed_host_findings_2026-10-23.jsonl | Provider managed-host findings, colocated estates ({info['feed']:,} rows) | Provider feed, current state |
| colocation_host_inventory_2026-10-23.xlsx | Colocated host inventory ({info['inv']:,} hosts) | Provider inventory export |
| colocation_capacity_register_2026-10.xlsx | Serving capacity and drain cycles | Capacity planning extract |
| estate_transaction_forecast_2026.csv | Hourly transaction forecast ({info['fc']:,} rows) | Forecasting export |
| colocation_change_acknowledgements_may_oct_2026.xlsx | Provider change acknowledgements ({info['acks']:,} rows) | Provider record |
| cloud_patch_ticket_log_q3_2026.csv / cloud_q3_closed_findings.csv | Q3 cloud tickets and the findings each closed | Ticketing export |
| q3_2026_remediation_closeout.pdf | Q3 close-out, cloud estates only | Office report |
| vulnerability_management_standard.docx | The office standard | Office document |
| sre_maintenance_standard.pdf | SRE maintenance standard | SRE document |
| colocation_service_schedule_and_window_calendar.pdf | Provider schedule and November calendar | Provider document |
| november_window_calendar.csv | November windows booked to the office, with drain cycles | Provider booking export |
| november_round_planning_thread.eml | Planning thread | Mail export |
| crew_deployment_log_2026.csv | Crew deployment log ({info['dep']:,} rows) | Deployment tooling export |
| cloud_asset_register_2026-10-23.csv | Cloud asset register ({info['asset']:,} rows) | Asset inventory export |
| scanner_coverage_2026-10.csv | Scanner coverage ({info['cov']:,} rows) | Scanner coverage export |
| vendor_advisory_feed.json | Vendor advisory first-release and revision dates | Advisory feed |
| drain_orchestrator_config.yaml | Provider drain orchestrator config | Provider configuration |
| superseded_cvss_band_allocation_memo.pdf | The prior CVSS-band allocation, superseded | Office document, retained for reference |
| warehouse_data_dictionary.md | Field meanings | Data platform document |

## Notes
- Amounts are whole counts of exposures. Times local to Madrid unless a file says otherwise.
- Nothing in the folder has been edited after extraction.
"""
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(t)
