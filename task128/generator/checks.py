"""task128 assertions: input gates, the ladder (every rung computed from the pack), the estimator family inside every
graded bin, the semantic fork rows SF1 to SF12, the asks' devices and stops, the separation line, consistency
between files and the realism sweep. Run by build.py after every build; exits on the first failure."""
import datetime as dt
import json
import os
import re
import shutil
import tempfile
import zipfile

import numpy as np
import pandas as pd
import pyarrow.parquet as pq

import calc as C
import docs as DC
import pack as P
import world as W

N_PASS = [0]


def ok(cond, label, detail=""):
    if not cond:
        raise SystemExit("CHECK FAILED: %s %s" % (label, detail))
    N_PASS[0] += 1
    print("  ok  %-62s %s" % (label, detail))


def demean(v, g):
    return v - pd.Series(v).groupby(g).transform("mean").values


def lpm(df, t, fe=None, sc=True):
    y = df.roll.astype(float).values
    x = df[t].astype(float).values
    s = ((df.BEH_SCORE.astype(float) - 650) / 50).values
    if fe is not None:
        g = df[fe].values
        y, x, s = demean(y, g), demean(x, g), demean(s, g)
    X = np.c_[np.ones(len(y)), x, s] if sc else np.c_[np.ones(len(y)), x]
    return float(np.linalg.lstsq(X, y, rcond=None)[0][1])


def iv_fe(df):
    g = df.day.values
    y = demean(df.roll.astype(float).values, g)
    d = demean(df.dialled.astype(float).values, g)
    z = demean((~df.left).astype(float).values, g)
    return float((z @ y) / (z @ d))


def day_weighted(ov):
    a = ov.groupby(["day", "left"]).roll.agg(["sum", "size"]).unstack("left")
    rb, ra = a[("sum", True)] / a[("size", True)], a[("sum", False)] / a[("size", False)]
    nb, na = a[("size", True)], a[("size", False)]
    d = rb - ra
    h = nb * na / (nb + na)
    return dict(std=float((nb * d).sum() / nb.sum()), byn=float(((nb + na) * d).sum() / (nb + na).sum()),
                byh=float((h * d).sum() / h.sum()))


def family(L, c):
    """Every reasonable same-list estimator of the rolls the new team would have prevented at a centre."""
    Y = L[(L.day >= C.YEAR_FROM) & (L.day <= C.YEAR_TO) & (L.CENTRE_CD == c)]
    over = sorted(Y[Y.left].day.unique())
    ov = Y[Y.day.isin(over)].copy()
    ov["reached_d"] = ~ov.left
    s = float(ov[~ov.left].skipped.mean())
    nb = ov[ov.left].groupby("day").size()
    dials = float(np.minimum((1 - s) * nb, C.NEW_TEAM).sum())
    dw = day_weighted(ov)
    itt = {"fe": -lpm(ov, "reached_d", "day"), "fe0": -lpm(ov, "reached_d", "day", sc=False), "std": dw["std"],
           "byn": dw["byn"], "byh": dw["byh"], "pool": -lpm(ov, "reached_d"), "pool0": -lpm(ov, "reached_d", sc=False)}
    out = {k: v / (1 - s) * dials for k, v in itt.items()}
    out["iv"] = -iv_fe(ov) * dials
    return out, dict(s=s, dials=dials, nb=int(nb.sum()), ov=ov, over=over)


def rungs(L, H, c):
    """The wrong sizings, each a complete pipeline from the pack."""
    Y = L[(L.day >= C.YEAR_FROM) & (L.day <= C.YEAR_TO) & (L.CENTRE_CD == c)]
    fam, info = family(L, c)
    ov, s, dials = info["ov"], info["s"], info["dials"]
    r = {}
    o2 = ov[~ov.skipped]
    r["R2"] = -lpm(o2, "dialled", "day") * dials
    r["R2_std"] = day_weighted(o2.assign(left=~o2.dialled))["std"] * dials
    r["R2_pool"] = -lpm(o2, "dialled") * dials
    r["R2b"] = -lpm(ov, "dialled", "day") * dials
    Yr = Y.assign(reached_d=~Y.left)
    r["R3"] = -lpm(Yr, "reached_d") / (1 - s) * dials
    r["R3_0"] = -lpm(Yr, "reached_d", sc=False) / (1 - s) * dials
    r["R1"] = -lpm(Y, "dialled") * dials
    itt = day_weighted(ov)["std"]
    r["R4"] = itt * info["nb"]
    r["noiv"] = itt * dials
    r["iv_all_left"] = itt / (1 - s) * info["nb"]
    # R0: every 30-day entry, listed or held, called or not
    held = H[H.HOLD_CD.notna() & (H.HOLD_CD != "")]
    return r, fam, info


def run(target, golden, world=None, meta_path=None):
    print("checks: task128")
    X = C.everything(target)
    L, H = X["L"], X["H"]
    S = X["size"]
    files = sorted(os.listdir(target))
    meta = json.load(open(meta_path))

    # ---- input gates
    ok(len(files) >= 10, "G1 ten or more files", str(len(files)))
    fmts = {f.rsplit(".", 1)[1] for f in files}
    ok(len(fmts) >= 3, "G2 three or more formats", ",".join(sorted(fmts)))
    ok(len(L) >= 25000, "G3 a file of 25,000 rows or more", "lists %d" % len(L))
    ok(len(meta["distractor_files"]) >= 2 and all(d in files for d in meta["distractor_files"]),
       "G4 two distractors named in metadata", ", ".join(meta["distractor_files"]))
    ok(not any("distractor" in f.lower() or "trap" in f.lower() or "golden" in f.lower() for f in files),
       "G5 no file name broadcasts the design")

    # ---- the over-capacity days and the new team's binding
    Y = L[(L.day >= C.YEAR_FROM) & (L.day <= C.YEAR_TO)]
    for c in C.CENTRES:
        D = Y[Y.CENTRE_CD == c]
        over = set(D[D.left].day)
        camp = set(D[D.CYCLE_CD.isin(["04", "13"])].day)
        ok(over == camp and len(over) == 24, "M1 %s over-capacity days are the 24 campaign-cycle days" % c,
           "%d days" % len(over))
    ok(S["LON"]["bind_lo"] >= 1.5, "M2 London's diallable left exceeds 400 by half on every day", "%.2f" % S["LON"]["bind_lo"])
    ok(max(S[c]["bind_hi"] for c in C.CENTRES if c != "LON") <= 0.8, "M3 no other centre's left reaches 0.8 of 400",
       "%.2f" % max(S[c]["bind_hi"] for c in C.CENTRES if c != "LON"))

    # ---- the golden and its family
    want = {"TBY": 100, "SUD": 100, "KGN": 200, "LON": 500, "WPG": 200}
    for c in C.CENTRES:
        ok(C.r100(S[c]["rolls"]) == want[c], "M4 %s golden rolls" % c, "%.1f -> %d" % (S[c]["rolls"], want[c]))
    ok(max(S[c]["rolls"] for c in C.CENTRES) < 0.75 * C.BAR, "M5 the largest centre is under three quarters of the line",
       "%.1f" % max(S[c]["rolls"] for c in C.CENTRES))
    ok(max(S, key=lambda c: S[c]["rolls"]) == "LON" and S["LON"]["rolls"] > 1.6 * max(S[c]["rolls"] for c in C.CENTRES if c != "LON"),
       "M6 London is the closest centre by a clear margin")
    allr = {}
    for c in C.CENTRES:
        r, fam, info = rungs(L, H, c)
        allr[c] = (r, fam, info)
        lo, hi = min(fam.values()), max(fam.values())
        b0 = want[c] - 50
        clear = min(lo - b0, b0 + 100 - hi)
        ok(clear >= 8, "SF11 %s every same-list estimator inside the golden's bin" % c,
           "%d to %d, clearance %.0f" % (lo, hi, clear))

    # ---- the ladder (London funded on every wrong rung)
    rL = allr["LON"][0]
    for k in ("R2", "R2_std", "R2_pool", "R2b", "R3", "R3_0", "R1", "R4"):
        ok(rL[k] >= 1.3 * C.BAR, "L1 %s funds London at 1.3 times the bar or more" % k, "%.0f" % rL[k])
    for k in ("R2", "R2b", "R3", "R1"):
        best = max(C.CENTRES, key=lambda c: allr[c][0][k])
        ok(best == "LON", "L2 %s names London" % k)
    # R0, the default path: every 30-day entry, listed or held; called if dialled on its list day
    held = H[H.HOLD_CD.notna() & (H.HOLD_CD != "")].copy()
    cd = H.groupby(["CYCLE_CD", H.CYCLE_DT.map(lambda d: (d.year, d.month))]).CYCLE_DT.first()
    def plus2(day, cyc):
        y, m = day.year, day.month + 2
        y, m = y + (m - 1) // 12, (m - 1) % 12 + 1
        return cd.get((cyc, (y, m)))
    held["day"] = held.CYCLE_DT
    held["out_dt"] = [plus2(d, c) for d, c in zip(held.day, held.CYCLE_CD)]
    late = H[H.DPD >= 90][["ACCT_NO", "CYCLE_DT"]].drop_duplicates().assign(roll=True)
    held = held.merge(late, left_on=["ACCT_NO", "out_dt"], right_on=["ACCT_NO", "CYCLE_DT"], how="left", suffixes=("", "_x"))
    held["roll"] = held.roll.fillna(False).astype(bool)
    acct_c = L.drop_duplicates("ACCT_NO").set_index("ACCT_NO").CENTRE_CD
    held["CENTRE_CD"] = held.ACCT_NO.map(acct_c)
    held = held[held.CENTRE_CD.notna()]
    held["dialled"] = False
    held["BEH_SCORE"] = 650.0
    r0 = {}
    for c in C.CENTRES:
        ent = pd.concat([Y[Y.CENTRE_CD == c][["roll", "dialled", "BEH_SCORE"]],
                         held[(held.day >= C.YEAR_FROM) & (held.day <= C.YEAR_TO) & (held.CENTRE_CD == c)][["roll", "dialled", "BEH_SCORE"]]])
        r0[c] = -lpm(ent, "dialled") * allr[c][2]["dials"]
    ok(r0["LON"] >= 1.3 * C.BAR and max(r0, key=r0.get) == "LON", "L3 R0 (the default path) funds London", "%.0f" % r0["LON"])
    # R0b: reached as the right party against not, among the dialled
    YL = Y[(Y.CENTRE_CD == "LON") & Y.dialled]
    pre = YL.day < X["upgrade"]["LON"]
    rpc = YL.RESULT_CD.isin(C.RPC_CODES) | ((YL.RESULT_CD == "AR") & pre)
    r0b = (YL[~rpc].roll.mean() - YL[rpc].roll.mean()) * allr["LON"][2]["dials"]
    ok(r0b >= 1.3 * C.BAR, "L4 R0b (contacted against not contacted) funds London", "%.0f" % r0b)
    # the grid: K (skipped kept with the reached), D (same list day), C (the new team's capacity)
    g = {"KDC": S["LON"]["rolls"], "KD-": rL["R4"], "K-C": rL["R3"], "-DC": rL["R2"]}
    ok(g["KDC"] < C.BAR and all(v >= 1.3 * C.BAR for k, v in g.items() if k != "KDC"),
       "L5 grid: only K, D and C together hold", ", ".join("%s %.0f" % kv for kv in g.items()))
    ok(C.r100(rL["noiv"]) != want["LON"] and rL["noiv"] < C.BAR, "L6 the effect left per listed record misses London's bin, same call",
       "%.0f" % rL["noiv"])

    # ---- semantic fork rows
    # SF1: ever 90 or more within 120 days of the list day gives the same rolls; at the next cycle date gives none
    late_any = H[H.DPD >= 90][["ACCT_NO", "CYCLE_DT"]]
    Yo = Y[["ACCT_NO", "day", "roll"]].merge(late_any, on="ACCT_NO", how="left")
    Yo["w"] = [(c is not None and not pd.isna(c) and 0 < (c - d).days <= 120) for c, d in zip(Yo.CYCLE_DT, Yo.day)]
    ever = Yo.groupby(["ACCT_NO", "day"]).w.any()
    gold = Y.set_index(["ACCT_NO", "day"]).roll
    ok(bool((ever.reindex(gold.index).fillna(False) == gold).all()), "SF1 ever 90 within 120 days selects the same rolls")
    # SF2: every record's list day is its cycle's collections day in the bank's calendar
    bad = sum(1 for d, cy in zip(L.day, L.CYCLE_CD) if W.collections_day(d.year, d.month, int(cy)) != d)
    ok(bad == 0, "SF2 every list day is its cycle's collections day", "0 of %d off" % len(L))
    ok(Y.day.min() == dt.date(2025, 7, 2) and Y.day.max() == dt.date(2026, 6, 25), "SF2 the year's list days run 2 July 2025 to 25 June 2026")
    # CP-14 6.2's "latest year of lists whose rolls are known" at the pull: August 2025 to July 2026
    alt = C.sizing(L, dt.date(2025, 8, 1), dt.date(2026, 7, 31))
    ok(all(C.r100(alt[c]["rolls"]) == C.r100(S[c]["rolls"]) for c in C.CENTRES)
       and max(alt[c]["rolls"] for c in C.CENTRES) < C.BAR,
       "SF2 the policy's latest year of known rolls (August 2025 to July 2026) keeps every bin and the call",
       ", ".join("%s %.0f" % (c, alt[c]["rolls"]) for c in C.CENTRES))
    # SF3, SF4: the existing teams' prevented rolls, or the centre's own capacity, move London off its bin
    own = S["LON"]["per_dial"] * W.CAP[3] * 24
    ok(C.r100(own) != want["LON"], "SF4 London sized at its own team's capacity leaves the bin", "%.0f" % own)
    exist = S["LON"]["per_dial"] * Y[(Y.CENTRE_CD == "LON")].dialled.sum()
    ok(C.r100(exist) != want["LON"], "SF3 the existing team's prevented rolls are another quantity", "%.0f" % exist)
    # SF5: each account sits at one centre
    ok(int(L.groupby("ACCT_NO").CENTRE_CD.nunique().max()) == 1, "SF5 every account is worked at one centre")
    # SF6: a new team that dialled the skipped records too moves Winnipeg off its bin; the call holds
    alt = {c: allr[c][0]["iv_all_left"] if c != "LON" else S["LON"]["rolls"] for c in C.CENTRES}
    ok(max(alt.values()) < C.BAR, "SF6 dialling the skipped records too keeps the call", "max %.0f" % max(alt.values()))
    # SF7: rounding the effect first keeps every bin
    for c in C.CENTRES:
        pr = round(100 * S[c]["per_dial"], 1) / 100 * S[c]["dials"]
        ok(C.r100(pr) == want[c], "SF7 %s effect rounded first keeps the bin" % c, "%.1f" % pr)
    # SF8: no test account in the lists or the history
    calls = pq.read_table(os.path.join(target, P.FILES["calls"])).to_pandas()
    qa = calls[calls.CAMPAIGN_ID.str.startswith("QA")]
    ok(len(qa) > 0 and not set(qa.ACCT_NO) & set(L.ACCT_NO) and not set(qa.ACCT_NO) & set(H.ACCT_NO),
       "SF8 the QA test accounts are in no list and no history", "%d test calls" % len(qa))
    # SF9: one record per account (first listing in the year) keeps the call
    first = Y.sort_values("day").drop_duplicates("ACCT_NO")
    Lf = L[L.index.isin(first.index) | ~((L.day >= C.YEAR_FROM) & (L.day <= C.YEAR_TO))]
    alt9 = {c: family(Lf, c) for c in C.CENTRES}
    mx = max(v[0]["std"] for v in alt9.values())
    ok(mx < C.BAR, "SF9 one record per account keeps the call", "max %.0f" % mx)
    # SF10, SF12 covered by M5, M2 and M3

    # ---- the asks
    K = X["contacts"]
    rpc_t = {"TBY": 50.4, "SUD": 54.6, "KGN": 57.2, "LON": 49.3, "WPG": 52.8}
    ptp_t = {"TBY": 41.2, "SUD": 36.4, "KGN": 39.1, "LON": 37.6, "WPG": 40.3}
    for c in C.CENTRES:
        ok(C.r1(K[c]["rpc_rate"]) == rpc_t[c] and abs(K[c]["rpc_rate"] - rpc_t[c]) <= 0.03,
           "A1 %s contacts per 100 dialled mid-bin" % c, "%.3f" % K[c]["rpc_rate"])
        ok(C.r1(K[c]["ptp_rate"]) == ptp_t[c] and abs(K[c]["ptp_rate"] - ptp_t[c]) <= 0.03,
           "A2 %s promises per 100 contacts mid-bin" % c, "%.3f" % K[c]["ptp_rate"])
    for c in C.CENTRES:
        D = Y[(Y.CENTRE_CD == c) & Y.dialled]
        nat = D.RESULT_CD.isin(C.RPC_CODES)
        over = D.RESULT_CD.isin(C.RPC_CODES | {"AR"})
        ptp = (D.RESULT_CD == "PTP").sum()
        r_nat, p_nat = 100 * nat.mean(), 100 * ptp / nat.sum()
        r_ov, p_ov = 100 * over.mean(), 100 * ptp / over.sum()
        ok(C.r1(r_nat) != rpc_t[c] and C.r1(p_nat) != ptp_t[c], "A3 %s the current code list read back misses both" % c,
           "%.1f / %.1f" % (r_nat, p_nat))
        ok(C.r1(r_ov) != rpc_t[c] and C.r1(p_ov) != ptp_t[c], "A4 %s AR a contact throughout misses both" % c,
           "%.1f / %.1f" % (r_ov, p_ov))
        one = D.day < W.UPGRADE["TBY"]
        rpc1 = D.RESULT_CD.isin(C.RPC_CODES) | ((D.RESULT_CD == "AR") & one)
        if c == "WPG":
            ok(C.r1(100 * rpc1.mean()) != rpc_t[c], "HZ1 Winnipeg on Ontario's upgrade day misses", "%.1f" % (100 * rpc1.mean()))
        dc = calls[(calls.CENTRE_CD == c) & (calls.LIST_DT >= C.YEAR_FROM) & (calls.LIST_DT <= C.YEAR_TO)]
        tc = dc[dc.CAMPAIGN_ID.str.startswith("QA")]
        rq = 100 * (rpc1.sum() if c != "WPG" else (D.RESULT_CD.isin(C.RPC_CODES) | ((D.RESULT_CD == "AR") & (D.day < W.UPGRADE["WPG"]))).sum())
        with_qa = (rq + 100 * tc.RESULT_CD.isin(C.RPC_CODES).sum()) / (len(D) + tc.ACCT_NO.size)
        ok(C.r1(with_qa) != rpc_t[c], "HZ2 %s the QA test calls counted miss" % c, "%.1f" % with_qa)
    # the device's evidence: old AR carries an agent and minutes of talk, new AR neither
    cd_ = calls[~calls.CAMPAIGN_ID.str.startswith("QA")].copy()
    cd_["pre"] = [d < W.UPGRADE[c] for d, c in zip(cd_.LIST_DT, cd_.CENTRE_CD)]
    oa, na = cd_[(cd_.RESULT_CD == "AR") & cd_.pre], cd_[(cd_.RESULT_CD == "AR") & ~cd_.pre]
    ok(oa.AGENT_ID.notna().all() and oa.TALK_SEC.min() >= 300 and na.AGENT_ID.isna().all() and na.TALK_SEC.max() < 60,
       "A5 the AR switch shows only across columns", "old AR %d calls, new AR %d" % (len(oa), len(na)))
    ok(not (cd_.pre & (cd_.RESULT_CD == "ARR")).any() and not (~cd_.pre & (cd_.RESULT_CD == "AM")).any(),
       "A6 ARR appears only after, AM only before each centre's upgrade")
    B = X["balances"]
    bal_t = {"TBY": 7000000, "SUD": 10100000, "KGN": 14300000, "LON": 15500000, "WPG": 16000000}
    for c in C.CENTRES:
        ok(C.r100k(B[c]["cad"]) == bal_t[c] and abs(B[c]["cad"] - bal_t[c]) <= 15000, "B1 %s balances mid-bin" % c,
           "%.0f" % B[c]["cad"])
        ok(C.r100k(B[c]["as_cad"]) != bal_t[c], "B2 %s US-dollar balances read as CAD miss" % c, "%.0f" % B[c]["as_cad"])
    # FX conventions: month-average and year-average rates stay in each bin
    A = pd.read_csv(os.path.join(target, P.FILES["arrangements"]), dtype={"ACCT_NO": str})
    A["setup"] = pd.to_datetime(A.SETUP_DT).dt.date
    A = A[(A.setup >= C.YEAR_FROM) & (A.setup <= C.YEAR_TO)]
    R = pd.read_csv(os.path.join(target, P.FILES["rates"]))
    R["d"] = pd.to_datetime(R.RATE_DT).dt.date
    prod = H.drop_duplicates("ACCT_NO").set_index("ACCT_NO").PRODUCT_CD
    A["usd"] = A.ACCT_NO.map(prod).eq("VUSD")
    mavg = R.groupby(R.d.map(lambda d: (d.year, d.month))).BOOKING_RATE.mean()
    yavg = R[(R.d >= C.YEAR_FROM) & (R.d <= C.YEAR_TO)].BOOKING_RATE.mean()
    for c in C.CENTRES:
        a = A[A.CENTRE_CD == c]
        m = np.where(a.usd, a.ARR_BAL * a.setup.map(lambda d: mavg[(d.year, d.month)]), a.ARR_BAL).sum()
        yv = np.where(a.usd, a.ARR_BAL * yavg, a.ARR_BAL).sum()
        dropped = a[~a.usd].ARR_BAL.sum()
        ok(C.r100k(m) == bal_t[c] and C.r100k(yv) == bal_t[c], "B3 %s month- and year-average rates keep the bin" % c)
        ok(C.r100k(dropped) != bal_t[c], "B4 %s US-dollar arrangements dropped miss" % c)

    # ---- the separation line
    ar = L[L.RESULT_CD == "AR"]
    ok((ar.ATTEMPTS >= 1).all(), "S1 AR is a dial result before and after the upgrade", "%d records" % len(ar))
    ok(((L.RESULT_CD == "") == (L.ATTEMPTS == 0) & ~L.RESULT_CD.isin(["DNC", "NDN"])).all(),
       "S2 blank results are exactly the records with no attempt that were not skipped")
    ok(int((L[L.RESULT_CD.isin(["DNC", "NDN"])].ATTEMPTS > 0).sum()) == 0, "S3 skipped records carry no attempt")

    # ---- consistency between files
    wl = pd.read_csv(os.path.join(target, P.FILES["worklog"]), dtype=str, keep_default_na=False)
    wl["day"] = pd.to_datetime(wl.WORK_DT).dt.date
    dial_wl = wl.groupby(["CENTRE_CD", "day"]).ES30_ACCOUNTS_DIALLED.apply(lambda s: s.astype(int).sum())
    dial_l = L.groupby(["CENTRE_CD", "day"]).dialled.sum()
    ok(bool((dial_wl.reindex(dial_l.index) == dial_l).all()), "X1 the work log's dials tie to the lists, pods summed")
    load_wl = wl[wl.ES30_RECORDS_LOADED != ""].groupby(["CENTRE_CD", "day"]).ES30_RECORDS_LOADED.apply(lambda s: s.astype(int).sum())
    ok(bool((load_wl.reindex(dial_l.index) == L.groupby(["CENTRE_CD", "day"]).size()).all()), "X2 records loaded tie to the lists")
    ent = H.set_index(["ACCT_NO", "CYCLE_DT"]).DPD
    li = L.set_index(["ACCT_NO", "day"]).DPD
    ok(bool((ent.reindex(li.index) == li).all()), "X3 every list record has its entry row at its list day, same DPD")
    ok(bool(set(A.ACCT_NO) <= set(L.ACCT_NO)), "X4 every arrangement's account was on a list")
    ok(int((calls.CALL_START.dt.hour < 8).sum() + (calls.CALL_START.dt.hour >= 20).sum()) == 0, "X5 every call between 08:00 and 20:00")
    ok(int(L.ATTEMPTS.max()) <= 3, "X6 no record dialled more than three times")
    vers = wl.groupby("CENTRE_CD").apply(lambda g: g.sort_values("day").DIALLER_VER.ne(g.sort_values("day").DIALLER_VER.shift()).sum())
    ok(bool((vers == 2).all()), "X7 each centre's dialler version changes once")
    ok(H.CYCLE_DT.max() <= dt.date(2026, 10, 8) and L.day.max() <= W.PULL, "X8 nothing dated after the pull")

    # ---- realism sweep
    for k in ("left", "dials"):
        vals = [round(S[c][k]) for c in C.CENTRES]
        ok(len(set(vals)) == 5, "R1 no two centres share a %s count" % k, str(vals))
    ok(all(round(S[c]["rolls"]) % 100 != 0 for c in C.CENTRES), "R2 no unrounded figure lands on a round hundred")
    ok(all(abs(B[c]["cad"] - round(B[c]["cad"], -5)) > 1000 for c in C.CENTRES), "R3 no balance total near a round figure")
    txt = ""
    for f in files:
        if f.endswith((".docx", ".xlsx")):
            with zipfile.ZipFile(os.path.join(target, f)) as z:
                txt += " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml"))
    ok("python-docx" not in txt and "openpyxl" not in txt.lower(), "R4 office containers scrubbed of the writing library")
    pdf = open(os.path.join(target, DC.PROTOCOL), "rb").read()
    ok(b"ReportLab" not in pdf and b"reportlab" not in pdf, "R5 the PDF carries no library banner")
    people = re.findall(r"(Elizabeth Marquez|Gabriel Ross|Kathryn Gray)", txt)
    ok(len(people) >= 2, "R6 the documents' people are the card's", ", ".join(sorted(set(people))))
    alltext = txt + pdf.decode("latin-1", "ignore")
    ok("—" not in alltext and "–" not in alltext, "R7 no em or en dash in the documents")
    stamps = {os.path.getmtime(os.path.join(target, f)) for f in files}
    ok(len(stamps) == 1, "R8 every file carries the export stamp")
    # every field the reference lists under a file's heading is a column of that file
    import docx
    ref = docx.Document(os.path.join(target, DC.FIELDS))
    cur, missing, seen = None, [], 0
    for el in ref.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[1]
        if tag == "p":
            t = "".join(x.text or "" for x in el.iter() if x.tag.endswith("}t")).strip()
            sty = el.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pStyle")
            if sty is not None and "Heading" in sty.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", ""):
                cur = t if t in files and t.endswith((".csv", ".parquet")) else None
        elif tag == "tbl" and cur:
            if cur.endswith(".csv"):
                cols = set(pd.read_csv(os.path.join(target, cur), nrows=0).columns)
            else:
                cols = set(pq.ParquetFile(os.path.join(target, cur)).schema_arrow.names)
            for tr in el.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tr"):
                cell = tr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tc")
                name = "".join(x.text or "" for x in cell.iter() if x.tag.endswith("}t")).strip()
                for n in name.split(", "):
                    if re.fullmatch(r"[A-Z][A-Z0-9_]+", n):
                        seen += 1
                        if n not in cols:
                            missing.append("%s.%s" % (cur, n))
    ok(seen >= 15 and not missing, "R9 every field the reference lists is a column of its file",
       "%d fields; missing %s" % (seen, missing))
    reg = pd.read_csv(os.path.join(target, DC.MANIFEST), dtype=str)
    note = reg.loc[reg.OBJECT_NAME == P.FILES["lists"], "SCOPE_NOTE"].iloc[0]
    span = "%d %s to %d %s" % (L.day.min().day, L.day.min().strftime("%b %Y"), L.day.max().day, L.day.max().strftime("%b %Y"))
    ok(span in note, "R10 the register's list-day span is the export's", note)

    # ---- distractors: deleting them changes nothing
    tmp = tempfile.mkdtemp()
    try:
        for f in files:
            if f not in meta["distractor_files"]:
                shutil.copy(os.path.join(target, f), tmp)
        X2 = C.everything(tmp)
        same = all(abs(X2["size"][c]["rolls"] - S[c]["rolls"]) < 1e-9 and X2["contacts"][c] == K[c] and
                   abs(X2["balances"][c]["cad"] - B[c]["cad"]) < 1e-6 for c in C.CENTRES)
        ok(same, "D1 the golden and every ask unchanged with the distractors deleted")
    finally:
        shutil.rmtree(tmp)

    # ---- goldens agree with figures.json
    fig = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "figures.json")))
    gc = pd.read_csv(os.path.join(golden, "centre_call_effects_2025_2026.csv"))
    for r in gc.itertuples(index=False):
        ok(str(int(r.ROLLS_PREVENTED_JUL25_JUN26)) == fig["rolls"][r.CENTRE].replace(",", "") and
           "%.1f" % r.RPC_PER_100_DIALLED == fig["rpc_per_100"][r.CENTRE] and
           "%.1f" % r.PTP_PER_100_RPC == fig["ptp_per_100_rpc"][r.CENTRE] and
           "{:,}".format(int(r.ARR_BAL_PLACED_CAD)) == fig["arr_bal_cad"][r.CENTRE],
           "F1 %s CSV row matches figures.json" % r.CENTRE)
    import docx
    memo = "\n".join(p.text for p in docx.Document(os.path.join(golden, "sixth_team_decision_2027.docx")).paragraphs)
    ok("return budget line C-01" in memo and fig["best_rolls"] in memo and fig["best_short"] in memo, "F2 the memo states the call and London's gap")
    print("checks: %d passed" % N_PASS[0])
