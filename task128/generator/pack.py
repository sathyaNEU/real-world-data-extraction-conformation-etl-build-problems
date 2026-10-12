"""Write the evidence pack: the dialler's list export, its call detail, the card platform's delinquency history, the workforce and dialler log, the arrangement extract, the treasury's booking rates and the
two files the teams' managers keep beside them."""
import csv
import datetime as dt
import io

import numpy as np
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq

import world as W

FILES = dict(
    lists="ES30_DIALLER_LISTS_2025-01_2026-09.csv",
    calls="ES30_CALL_DETAIL_2025-07_2026-06.parquet",
    history="CARD_DELINQUENCY_HISTORY_2025-01_2026-10.parquet",
    worklog="ES_WORKFORCE_DIALLER_LOG_2025-01_2026-09.csv",
    arrangements="ES30_PAYMENT_ARRANGEMENTS_2025-01_2026-09.csv",
    rates="TREASURY_BOOKING_RATES_USDCAD_2025-01_2026-10.csv",
    ls60="LS60_DAILY_LIST_SUMMARY_2025-01_2026-09.csv",
    headcount="Collections_Workforce_Headcount_2025Q1-2026Q3.xlsx",
)
DISTRACTORS = ["ls60", "headcount"]


def write_text(path, text):
    with open(path, "w", encoding="utf-8", newline="") as fh:
        fh.write(text)


def _csv(rows, header):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(header)
    for r in rows:
        w.writerow(["" if v is None else v for v in r])
    return buf.getvalue()


def d8(d):
    return d.strftime("%Y-%m-%d")


def lists(E, acctno, path):
    L = E[E.onlist].sort_values(["day", "centre", "seq"])
    rows = []
    for e in L.itertuples(index=False):
        cc = W.CODES[e.centre]
        rows.append((d8(e.day), cc, int(e.seq), acctno[e.acct], "%02d" % e.cycle, int(e.dpd0),
                     int(e.score), int(e.attempts), e.code))
    write_text(path, _csv(rows, ["LIST_DT", "CENTRE_CD", "LIST_SEQ", "ACCT_NO", "CYCLE_CD", "DPD", "BEH_SCORE",
                                 "ATTEMPTS", "RESULT_CD"]))
    return len(rows)


def _ts(day, minute, tz_shift=0):
    t = dt.datetime.combine(day, dt.time(8, 0)) + dt.timedelta(minutes=float(minute))
    return t.strftime("%Y-%m-%d %H:%M:%S")


def calls(R, T, acctno, rng, path):
    """Attempt rows from the dialler's call detail: the old dialler numbered its calls, the new one keys them."""
    R = R.copy()
    R["acct_no"] = acctno[R.acct.values]
    R["campaign"] = ["ES30_%s" % W.CODES[c] for c in R.centre.values]
    R["cc"] = [W.CODES[c] for c in R.centre.values]
    T = T.copy()
    T["campaign"] = ["QA_%s" % c for c in T.centre.values]
    T["cc"] = T.centre
    T["attempt"] = 1
    T["pre"] = [d < W.UPGRADE[c] for d, c in zip(T.day.values, T.centre.values)]
    cols = ["day", "minute", "cc", "campaign", "acct_no", "attempt", "code", "talk", "agent", "pre"]
    X = pd.concat([R[cols], T[cols]], ignore_index=True)
    X["ts"] = [_ts(d, m) for d, m in zip(X.day.values, X.minute.values)]
    X = X.sort_values(["ts", "cc", "acct_no", "attempt"]).reset_index(drop=True)
    old_n = 7720184400
    rows = []
    for x in X.itertuples(index=False):
        if x.pre:
            old_n += int(rng.integers(1, 4))
            cid = "%d" % old_n
        else:
            cid = "CX%010X" % int(rng.integers(0, 2 ** 40))
        rows.append((cid, x.campaign, x.cc, x.day, x.acct_no, int(x.attempt), x.ts, x.agent, x.code, int(x.talk)))
    cols = list(zip(*rows))
    t = pa.table({"CALL_ID": pa.array(cols[0]), "CAMPAIGN_ID": pa.array(cols[1]), "CENTRE_CD": pa.array(cols[2]),
                  "LIST_DT": pa.array(cols[3], type=pa.date32()), "ACCT_NO": pa.array(cols[4]),
                  "ATTEMPT_NO": pa.array(cols[5], type=pa.int16()),
                  "CALL_START": pa.array([dt.datetime.strptime(v, "%Y-%m-%d %H:%M:%S") for v in cols[6]], type=pa.timestamp("s")),
                  "AGENT_ID": pa.array([v if v else None for v in cols[7]], type=pa.string()),
                  "RESULT_CD": pa.array(cols[8]), "TALK_SEC": pa.array(cols[9], type=pa.int32())})
    pq.write_table(t, path, compression="zstd", compression_level=9, row_group_size=200000, use_dictionary=True,
                   version="2.6", data_page_version="1.0", store_schema=False)
    return len(rows)


def history(H, A, acctno, path):
    X = H.sort_values(["cdt", "cycle", "acct"]).reset_index(drop=True)
    t = pa.table({"ACCT_NO": pa.array(acctno[X.acct.values].astype(str)),
                  "PRODUCT_CD": pa.array(A["product"].values[X.acct.values].astype(str)),
                  "CYCLE_CD": pa.array(["%02d" % c for c in X.cycle.values]),
                  "CYCLE_DT": pa.array(list(X.cdt.values), type=pa.date32()),
                  "DPD": pa.array(X.dpd.values.astype(np.int32)),
                  "BAL_CAD": pa.array(X.bal_cad.values.astype(np.float64)),
                  "HOLD_CD": pa.array([s if s else None for s in X.hold.values], type=pa.string()),
                  "STATUS_CD": pa.array([s if s else None for s in X.status.values], type=pa.string())})
    pq.write_table(t, path, compression="zstd", compression_level=9, row_group_size=250000, write_statistics=True,
                   use_dictionary=True, version="2.6", data_page_version="1.0", store_schema=False)
    return len(X)


def worklog(C, path):
    X = C.sort_values(["day", "centre", "team"])
    rows = [(d8(r.day), r.centre, r.team, r.ver, int(r.seats), "%.2f" % r.hours,
             "" if r.loaded is None or (isinstance(r.loaded, float) and np.isnan(r.loaded)) else int(r.loaded), int(r.dialled))
            for r in X.itertuples(index=False)]
    write_text(path, _csv(rows, ["WORK_DT", "CENTRE_CD", "TEAM_CD", "DIALLER_VER", "SEATS_STAFFED", "AGENT_HRS",
                                 "ES30_RECORDS_LOADED", "ES30_ACCOUNTS_DIALLED"]))
    return len(rows)


def arrangements(R, acctno, path):
    rows = [(r.ref, acctno[r.acct], W.CODES[r.centre], d8(r.setup), r.agent, "%.2f" % r.arr_bal, int(r.instalments),
             r.freq, d8(r.first_due), r.status) for r in R.itertuples(index=False)]
    write_text(path, _csv(rows, ["ARR_REF", "ACCT_NO", "CENTRE_CD", "SETUP_DT", "SETUP_AGENT", "ARR_BAL", "INSTALMENTS",
                                 "FREQUENCY", "FIRST_DUE_DT", "STATUS"]))
    return len(rows)


def rates(Rt, path):
    rows = [(d8(r.day), "USD", "CAD", "%.4f" % r.rate) for r in Rt.itertuples(index=False)]
    write_text(path, _csv(rows, ["RATE_DT", "FROM_CCY", "TO_CCY", "BOOKING_RATE"]))
    return len(rows)


def ls60(G, path):
    X = G.sort_values(["cdt", "centre"])
    rows = [(d8(r.cdt), W.CODES[r.centre], int(r.loaded), int(r.dialled), int(r.rpc), int(r.ptp)) for r in X.itertuples(index=False)]
    write_text(path, _csv(rows, ["LIST_DT", "CENTRE_CD", "LS60_RECORDS_LOADED", "LS60_ACCOUNTS_DIALLED", "LS60_RPC",
                                 "LS60_PTP"]))
    return len(rows)


def headcount(Hc, path):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Headcount by quarter"
    ws.append(["Collections workforce by centre, every collections role (30-day, 60-day and recoveries teams, "
               "team leaders and quality)"])
    ws.append(["Source: HR quarter-end extract; leavers include internal moves out of Collections"])
    ws.append([])
    hdr = ["Quarter", "Centre", "Headcount at quarter end", "Open vacancies", "Leavers in quarter", "Attrition %"]
    ws.append(hdr)
    for r in Hc.itertuples(index=False):
        ws.append([r.quarter, r.centre, int(r.headcount), int(r.vacancies), int(r.leavers), float(r.attrition_pct)])
    ws["A1"].font = Font(bold=True, size=11)
    ws["A2"].font = Font(italic=True, size=9, color="555555")
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCE6F1")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=5):
        row[5].number_format = "0.0"
    widths = [10, 13, 14, 11, 12, 10]
    for i, w in enumerate(widths):
        ws.column_dimensions["ABCDEF"[i]].width = w
    ws.freeze_panes = "A5"
    wb.save(path)
    return len(Hc)
