"""The golden deliverables, computed forward from the shipped files: centre_call_effects_2025_2026.csv (each
centre's rolls the new team would have prevented, contacts, promises and arrangement balances),
centre_call_effects_2025_2026.png (the five centres against the policy's funding line) and
sixth_team_decision_2027.docx (the decision), with generator/figures.json for the write-up."""
import json
import os

import calc as C

CSV = "centre_call_effects_2025_2026.csv"
PNG = "centre_call_effects_2025_2026.png"
DOCX = "sixth_team_decision_2027.docx"


def fmt(n):
    return "{:,}".format(int(n))


def figures(X):
    S, K, B = X["size"], X["contacts"], X["balances"]
    rolls = {C.NAMES[c]: C.r100(S[c]["rolls"]) for c in C.CENTRES}
    order = sorted(C.CENTRES, key=lambda c: (-S[c]["rolls"], C.NAMES[c]))
    best = order[0]
    f = dict(
        call="none",
        bar_rolls=fmt(C.BAR),
        new_team_dials=fmt(C.NEW_TEAM),
        best_centre=C.NAMES[best],
        best_rolls=fmt(rolls[C.NAMES[best]]),
        best_short=fmt(C.r100(C.BAR - S[best]["rolls"])),
        over_days=fmt(S[best]["over_days"]),
        rolls={C.NAMES[c]: fmt(rolls[C.NAMES[c]]) for c in C.CENTRES},
        order=[C.NAMES[c] for c in order],
        london_dials=fmt(S["LON"]["dials"]),
        london_per_100_dials="%.1f" % C.r1(100 * S["LON"]["per_dial"]),
        rpc_per_100={C.NAMES[c]: "%.1f" % C.r1(K[c]["rpc_rate"]) for c in C.CENTRES},
        ptp_per_100_rpc={C.NAMES[c]: "%.1f" % C.r1(K[c]["ptp_rate"]) for c in C.CENTRES},
        arr_bal_cad={C.NAMES[c]: fmt(C.r100k(B[c]["cad"])) for c in C.CENTRES},
        upgrade={C.NAMES[c]: X["upgrade"][c].isoformat() for c in C.CENTRES},
        roll_dpd="90",
        balance_rounding="100,000",
        per="100",
        policy_clauses=["6.3", "6.4", "6.5"],
        campaign_cycles=["04", "13"],
    )
    return f, order


def write_csv(path, X, f, order):
    import csv
    S, K, B = X["size"], X["contacts"], X["balances"]
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\r\n")
        w.writerow(["CENTRE", "ROLLS_PREVENTED_JUL25_JUN26", "MEETS_CP14_LINE", "RPC_PER_100_DIALLED",
                    "PTP_PER_100_RPC", "ARR_BAL_PLACED_CAD"])
        for c in order:
            w.writerow([C.NAMES[c], C.r100(S[c]["rolls"]), "N", "%.1f" % C.r1(K[c]["rpc_rate"]),
                        "%.1f" % C.r1(K[c]["ptp_rate"]), C.r100k(B[c]["cad"])])


def write_png(path, X, f, order):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "svg.hashsalt": "t128",
                         "text.parse_math": False})
    S = X["size"]
    names = [C.NAMES[c] for c in order]
    vals = [C.r100(S[c]["rolls"]) for c in order]
    fig, ax = plt.subplots(figsize=(8.2, 4.9), dpi=150)
    colors = ["#1f5f73" if i == 0 else "#9fb7bf" for i in range(len(order))]
    bars = ax.bar(names, vals, color=colors, width=0.62, zorder=2)
    ax.axhline(C.BAR, color="#b2432f", lw=1.6, ls="--", zorder=3)
    ax.text(len(names) - 0.5, C.BAR + 14, "CP-14 funding line: %s rolls" % fmt(C.BAR), ha="right", va="bottom",
            color="#b2432f", fontsize=9.5)
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width() / 2, v + 12, fmt(v), ha="center", va="bottom", fontsize=9.5)
    ax.set_ylim(0, 820)
    ax.set_ylabel("Rolls to 90 days past due prevented a year")
    ax.set_title("No centre's new team would reach the %s-roll funding line" % fmt(C.BAR), loc="left", fontsize=12, pad=26)
    ax.text(0, 1.035, "Rolls the new team's calls would have prevented on each centre's July 2025 to June 2026 "
            "lists, nearest 100", transform=ax.transAxes, fontsize=8.8, color="#444444")
    ax.spines[["top", "right"]].set_visible(False)
    ax.yaxis.grid(True, color="#e3e3e3", zorder=0)
    fig.text(0.01, 0.01, "Same-list comparison on the campaign cycles' collections days; new team at 400 accounts "
             "dialled a day. Source: dialler lists, card delinquency history.", fontsize=7.6, color="#555555")
    fig.tight_layout(rect=(0, 0.03, 1, 1))
    fig.savefig(path, dpi=150, metadata={"Software": None, "Creation Time": None})
    plt.close(fig)


def write_docx(path, X, f, order):
    import docx
    from docx.shared import Pt, RGBColor
    S, K, B = X["size"], X["contacts"], X["balances"]
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "Ferncliffe Trust Bank  |  Collections Operations  |  For the Executive Committee"
    sec.footer.paragraphs[0].text = "Internal. Case under Credit Policy CP-14 section 6, filed with the 2027 budget request."
    head = [("To", "Executive Committee"), ("From", "Elizabeth Marquez, Director, Collections Operations"),
            ("Date", "14 October 2026"), ("Re", "Budget line C-01: the 2027 early-stage team"),
            ("Data", "Collections MI extracts of 9 October 2026")]
    for k, v in head:
        p = d.add_paragraph()
        r = p.add_run("%s: " % k)
        r.bold = True
        p.add_run(v)
        p.paragraph_format.space_after = Pt(0)
    d.add_paragraph()
    t = d.add_paragraph()
    r = t.add_run("Recommendation: return budget line C-01. No centre's new team clears the CP-14 funding line of "
                  "%s rolls; London comes closest at about %s rolls a year, %s short." % (
                      f["bar_rolls"], f["best_rolls"], f["best_short"]))
    r.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    d.add_paragraph(
        "I sized each centre as CP-14 section 6 sets out, on its 30-day lists of July 2025 to June 2026. Our teams "
        "only run out of shift on the collections days of cycles 04 and 13, the two cycles that carry the 2024 "
        "student-card campaign; on those %s days the records left at the bottom of each list when the campaign closed "
        "are the accounts a second team would have called. Because the list loads in the card platform's extract "
        "order, they are like-for-like with the records above them on the same list, so the difference in rolls "
        "between the two is what a call does. The dialler only checks the do-not-call file and the number when it "
        "reaches a record, so the records it skipped stay with the records it reached, and the effect is scaled to "
        "the records it actually dialled." % f["over_days"])
    d.add_paragraph(
        "The new team is rated at %s accounts dialled a day (C-01). At four centres it would have dialled every "
        "diallable record their teams left; at London, where the student-card cycles leave well over %s a day, it "
        "would have dialled its %s each of those days, %s dials in the year at %s rolls prevented per 100 dials." % (
            f["new_team_dials"], f["new_team_dials"], f["new_team_dials"], f["london_dials"], f["london_per_100_dials"]))
    tbl = d.add_table(rows=1, cols=3)
    tbl.style = "Light List Accent 1"
    for i, h in enumerate(["Centre", "Rolls prevented a year (nearest 100)", "Against the %s line" % fmt(C.BAR)]):
        tbl.rows[0].cells[i].text = h
    for c in order:
        row = tbl.add_row().cells
        row[0].text = C.NAMES[c]
        row[1].text = fmt(C.r100(S[c]["rolls"]))
        row[2].text = "short by %s" % fmt(C.r100(C.BAR - S[c]["rolls"]))
    src = d.add_paragraph("Source: dialler list export and card delinquency history, list days July 2025 to June 2026; "
                          "rolls rounded to the nearest 100.")
    for run in src.runs:
        run.font.size = Pt(8.5)
        run.italic = True
    d.add_paragraph("What each team's calls achieved over the same year, for the committee's reference:")
    t2 = d.add_table(rows=1, cols=4)
    t2.style = "Light List Accent 1"
    for i, h in enumerate(["Centre", "Right-party contacts per 100 accounts dialled", "Promises to pay per 100 contacts",
                           "Balances placed on arrangements (CAD, nearest 100,000)"]):
        t2.rows[0].cells[i].text = h
    for c in order:
        row = t2.add_row().cells
        row[0].text = C.NAMES[c]
        row[1].text = "%.1f" % C.r1(K[c]["rpc_rate"])
        row[2].text = "%.1f" % C.r1(K[c]["ptp_rate"])
        row[3].text = fmt(C.r100k(B[c]["cad"]))
    src2 = d.add_paragraph("Source: dialler list export (July 2025 to June 2026 list days) and arrangement system "
                           "extract (arrangements set up in the same months).")
    for run in src2.runs:
        run.font.size = Pt(8.5)
        run.italic = True
    n = d.add_paragraph(
        "Notes. Contacts count an account each day it was dialled. Before each centre's move to the CX2 dialler (20 "
        "October 2025 in Ontario, 17 November 2025 in Winnipeg) the old dialler coded an arrangement AR, the code CX2 "
        "now uses for an answering machine; those calls are counted as contacts. Arrangement balances on US-dollar "
        "cards are converted at Treasury's booking rate on the set-up day.")
    for run in n.runs:
        run.font.size = Pt(9)
        run.italic = True
    d.save(path)


def main(target, out, figures_path):
    os.makedirs(out, exist_ok=True)
    X = C.everything(target)
    f, order = figures(X)
    write_csv(os.path.join(out, CSV), X, f, order)
    write_png(os.path.join(out, PNG), X, f, order)
    write_docx(os.path.join(out, DOCX), X, f, order)
    import postprocess as PP
    PP.normalize_ooxml(os.path.join(out, DOCX), "Elizabeth Marquez", "Elizabeth Marquez", "2026-10-13T20:14:00Z",
                       "2026-10-14T13:02:00Z", title="2027 early-stage team")
    stamp = 1791993720.0
    for name in os.listdir(out):
        os.utime(os.path.join(out, name), (stamp, stamp))
    with open(figures_path, "w") as fh:
        json.dump(f, fh, indent=1, sort_keys=True)
        fh.write("\n")
    return X, f
