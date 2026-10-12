"""Assertions on the built pack and goldens: the input gates; the three rules executed on the generator's clean truth
(R1 the field's raw material ships, R2 each Layer 1 item is neutral, R3a the easy route files the wrong premium, R3b the
decisive step alone files the golden, R3c no plant); every rung at its named wrong answer with its margin; no single
offset reaching the golden; the semantic forks SF1 to SF15; the rounding bins; each ask's stops outside the golden at
every sub-district and season; the documents' silence on the size of the change; the realism sweep; and the goldens
agreeing with figures.json."""
import collections
import csv
import datetime as dt
import json
import os
import re

import numpy as np
import openpyxl

import analysis as AN
import book as BK
import build as B
import docs as D
import pricing as P
import records as R
import world as W

N = [0]
HERE = os.path.dirname(os.path.abspath(__file__))


def ok(cond, what):
    N[0] += 1
    if not cond:
        raise AssertionError(what)


def _pdf_text(path):
    from pypdf import PdfReader
    return " ".join(pg.extract_text() for pg in PdfReader(path).pages).replace("\n", " ")


def _ws(text):
    return re.sub(r"\s+", " ", text)


def ask_a_stops(bl):
    """the cracker's and the lazy solver's routes to the sub-district cut"""
    ins = BK.schedule_2027(bl)
    out = {}
    lines = collections.defaultdict(lambda: [0, 0.0])
    blocks = collections.defaultdict(lambda: [0, 0.0])
    for b in ins:
        sd = BK.stored_district(b)
        for _ in b.varieties:
            lines[sd][0] += 1
            lines[sd][1] += b.area
        blocks[sd][0] += 1
        blocks[sd][1] += b.area
    out["schedule labels, variety lines summed"] = lines
    out["schedule labels, one row a block"] = blocks
    by_no = {}
    for b in ins:
        by_no.setdefault(b.no, b)
    one = collections.defaultdict(lambda: [0, 0.0])
    for b in by_no.values():
        one[b.district_new][0] += 1
        one[b.district_new][1] += b.area
    out["deduplicated on block number alone"] = one
    reg = collections.defaultdict(lambda: [0, 0.0])
    for b in bl:
        reg[b.district_new][0] += 1
        reg[b.district_new][1] += b.area
    out["every block in the co-operative's register"] = reg
    return {k: {sd: (v[sd][0], round(v[sd][1], 1)) for sd in W.SUB_DISTRICTS} for k, v in out.items()}


def ask_b_stops(st):
    reg, ledger = st["claims"], st["ledger"]
    out = {}
    first = {}
    for r in reg:
        first.setdefault(r["claim_no"], r)
    by = collections.defaultdict(lambda: [0, 0.0])
    for r in first.values():
        by[r["season"]][0] += 1
        by[r["season"]][1] += r["assessed_amount"]
    out["register claims, one amount a claim"] = by
    rows = collections.defaultdict(lambda: [0, 0.0])
    for r in reg:
        rows[r["season"]][0] += 1
        rows[r["season"]][1] += r["assessed_amount"]
    out["register rows summed"] = rows
    fy = collections.defaultdict(lambda: [set(), 0.0])
    for p in ledger:
        y = int(p["fin_year"][2:])
        lab = W.season_label(y - 1)
        fy[lab][0].add(p["claim_no"])
        fy[lab][1] += p["amount"]
    out["ledger by financial year"] = {k: [len(v[0]), v[1]] for k, v in fy.items()}
    labs = [W.season_label(s) for s in range(W.COVER_FIRST_SEASON, W.LAST_SEASON + 1)]
    return {k: {lab: (v[lab][0], P.r_dollar(v[lab][1])) if lab in v else (0, 0) for lab in labs} for k, v in out.items()}


def run(target, golden, st, meta_path):
    mc, fig = st["mc"], st["fig"]
    mf, gold, rungs, cv = st["mf"], mc["golden"], mc["rungs"], mc["conv"]
    figj = json.load(open(os.path.join(HERE, "figures.json")))
    ok(figj == json.loads(json.dumps(fig)), "figures.json is this build's figures")
    files = sorted(os.listdir(target))
    texts = {f: _ws(_pdf_text(os.path.join(target, f))) for f in files if f.endswith(".pdf")}
    texts.update({f: open(os.path.join(target, f), encoding="utf-8").read() for f in files if f.endswith((".md", ".json"))})
    manual, terms27, history = texts[D.MANUAL], texts[D.TERMS_27], texts[D.HISTORY]
    notes = texts[D.NOTES]

    # ---- input gates
    meta = json.load(open(meta_path))
    ok(len(files) == len(B.FILES) >= 10, "ten or more files")
    ok(len({f.rsplit(".", 1)[1] for f in files}) >= 3, "three or more formats")
    ok(meta["input_gates"]["largest_file_rows"] >= 25000, "a file of 25,000 rows or more")
    ok(len(meta["distractor_files"]) >= 2 and all(f in files for f in meta["distractor_files"]), "two distractors")
    vsrc = open(os.path.join(HERE, "verify.py")).read()
    for f in meta["distractor_files"]:
        ok("distract" not in f, "distractor not named as such")
        ok(f not in vsrc, "the solution never reads " + f)
    main_src = vsrc.split("def ask_a")[0]
    for f in (R.SCHEDULE, R.BLOCKS, R.CLAIMS, R.LEDGER):
        ok(f not in main_src, "separation: the main call's path never reads " + f)
    for f in (R.DAILY, R.REGISTER, R.PARALLEL):
        ok(f not in vsrc.split("def ask_a")[1], "separation: the asks never read " + f)

    # ---- R1: the field's raw material ships, the deciding quantity does not
    par = list(csv.DictReader(open(os.path.join(target, R.PARALLEL), encoding="utf-8")))
    ok(len(par) == len(st["pairs"]) > 10000, "R1 the paired readings ship day by day at seven sites")
    ok(list(par[0]) == ["site_no", "site_name", "obs_date", "max_glass_large_screen", "max_probe_small_shield"],
       "R1 the paired file carries the two readings and nothing derived")
    ok(len({r["site_no"] for r in par}) == 7, "R1 seven comparison sites")
    for f in files:
        if f.endswith(".csv"):
            head = open(os.path.join(target, f), encoding="utf-8").readline().lower()
            for w in ("diff", "adjust", "homog", "convert", "offset", "premium", "burn"):
                ok(w not in head, "R1 no shipped column holds a derived quantity (%s in %s)" % (w, f))
    for f, t in texts.items():
        ok(("$%d" % gold["premium"]) not in t and ("%d a hectare" % gold["premium"]) not in t,
           "R1 no shipped document states the premium (%s)" % f)
    ok("max_glass_large_screen" in vsrc and "groupby(\"max_glass_large_screen\")" in vsrc,
       "R1 the verifier's decisive computation reads the paired rows by reading")

    # ---- R2: each Layer 1 item is real and neutral
    daily = list(csv.DictReader(open(os.path.join(target, R.DAILY), encoding="utf-8")))
    keys = collections.Counter((r["obs_date"], r["element"]) for r in daily)
    dup = sorted({k[0] for k, v in keys.items() if v > 1})
    ok(dup and all(x.startswith("2012-07") for x in dup), "R2 the resent batch is July 2012 only")
    ok(all(W.season_of(dt.date.fromisoformat(x)) is None for x in dup), "R2 the resent batch is outside every season")
    prov = [r["obs_date"] for r in daily if r["quality"] == "P"]
    ok(prov and min(prov) > W.APPROVED_TO.isoformat(), "R2 provisional readings only after the latest approval run")
    ok(all(W.season_of(dt.date.fromisoformat(x)) is None for x in prov), "R2 provisional readings outside every season")
    reg_days = st["register_days"]
    probe_days = {dt.date.fromisoformat(r["obs_date"]) for r in daily if r["element"] == "TMAX" and r["quality"] == "A"}
    for s in P.SEASONS:
        for d in W.season_days(s):
            ok(d in (probe_days if d >= W.CHANGE else reg_days), "R2 no season day missing (%s)" % d)
    old_seasons = sum(1 for s in P.SEASONS if all(d in reg_days for d in W.season_days(s)))
    new_seasons = sum(1 for s in P.SEASONS if all(d in probe_days for d in W.season_days(s)))
    ok(old_seasons == 30 and new_seasons == 32, "R2 either station number alone carries 30 or 32 of 62 seasons (a "
                                                "visible gap, never a silent answer)")
    stations = json.load(open(os.path.join(target, R.STATIONS)))
    old = [x for x in stations["stations"] if x["station_no"] == W.OLD_ID][0]
    ok(old["successor"] == W.NEW_ID and old["closed"] == "1995-10-31", "R2 the station list joins the two numbers")
    ok(max(reg_days) < W.CHANGE, "R2 the register ends before the change (SF10)")
    for f in (R.SCHEDULE, R.BLOCKS, R.CLAIMS, R.LEDGER):
        ok(f not in main_src, "R2 the book's mess (lines, endorsements, duplicates, runs) touches only the asks: " + f)

    # ---- R3a, R3b: the easy route and the decisive step, on the clean truth
    ok(rungs["R1 mean difference"]["premium"] == rungs["R0 no conversion"]["premium"] != gold["premium"],
       "R3a the easy route (the mean conversion, every filed rule applied) files the raw record's premium")
    ok(rungs["R1 mean difference"]["premium"] / gold["premium"] - 1 >= 0.06, "R3a the easy route 6 per cent or more off")
    ok(rungs["R4 at each reading, pooled"]["premium"] == gold["premium"], "R3b the decisive step alone files the golden")
    for k in ("R2 summer mean difference", "R2 linear fit over all days", "R3 one offset, glass at 33 C or more",
              "R3 one offset, glass at 35 C or more"):
        ok(abs(rungs[k]["premium"] / gold["premium"] - 1) >= 0.02, "R3b %s leaves the premium 2 per cent or more off" % k)
    ok(rungs["R3 one offset, glass at 33 C or more"]["premium"] != rungs["R3 one offset, glass at 35 C or more"]["premium"],
       "the two hot offsets file different premiums")
    ok("%.1f" % rungs["R3 one offset, glass at 33 C or more"]["expected_severe"] != fig["expected_severe_days_2027_28"],
       "the 33 C offset misses the severe count")
    ok("%.1f" % rungs["R3 one offset, glass at 35 C or more"]["expected_heat"] != fig["expected_heat_days_2027_28"],
       "the 35 C offset misses the heat count")
    sweep = {round(u, 2): P.rung(mf, P.uniform(u))[2]["premium"] for u in np.arange(0.0, 2.001, 0.05)}
    ok(gold["premium"] not in sweep.values(), "no single offset from 0.0 C to 2.0 C reaches the golden")
    pooled = cv["pooled"]
    ok(pooled[34.5] >= 0.6 and pooled[34.0] <= 0.8 and pooled[39.0] >= 1.2 and pooled[38.5] <= 1.4,
       "the conversion moves 34.5 C over 35.0 C and 39.0 C over 40.0 C, and neither 34.0 C nor 38.5 C")
    # R3c: no plant
    for series, name in ((np.array([reg_days[d] for d in sorted(reg_days)]), "register"),
                         (np.array([float(mf["probe"][k]) for k, d in enumerate(mf["days"]) if d >= W.CHANGE]), "probe")):
        seen, rep = set(), 0
        for i in range(len(series) - 9):
            key = tuple(series[i:i + 10])
            rep += key in seen
            seen.add(key)
        ok(rep == 0, "R3c no repeated ten-day block in the %s series" % name)
    for f, t in texts.items():
        low = t.lower()
        for w in ("adjust", "homogen", "offset", "convert", "conversion", "difference", "warmer", "bias", "reads high",
                  "overstate", "inflat"):
            ok(w not in low, "R3c and silence: %s says nothing of the change's size (%r)" % (f, w))
        ok("\u2014" not in t and "\u2013" not in t, "no dash in " + f)

    # ---- semantic forks
    ok("is 35.0 C or more" in terms27, "SF1 heat day at 35.0 C or more")
    ok("a severe heat day is also a heat day" in terms27, "SF2 a severe heat day is also a heat day")
    ok("A season runs from 1 December to the last day of February" in manual, "SF3 the season")
    ok(sum(1 for s in P.SEASONS if len(W.season_days(s)) == 91) == sum(1 for s in P.SEASONS if (s + 1) % 4 == 0),
       "SF3 leap Februaries counted")
    ok("probe and shield not replaced" in history, "SF4 no instrument change since 1995")
    ok("from 1 December 1965" in manual, "SF5 the record from 1965/66")
    ok("on its own line" in manual and "least squares against the season's first year" in manual, "SF6 the lines")
    ok(not gold["floor_binds"], "SF7 no projected count below zero")
    ok("payout a hectare each projected season would make" in manual, "SF8 the burn is the average payout")
    ok("rounded once to the whole dollar" in manual, "SF9 rounded once")
    ok("became the station's official instrument" in history and "1995-10-31" == old["closed"],
       "SF10 the probe is the record from 1 November 1995")
    ok(rungs["R4 at each reading, Merrin Flat only"]["premium"] == gold["premium"], "SF11 Merrin Flat's pairs alone")
    ov = {d: (P.converter(cv["pooled"])(float(gg)) if d <= W.CHECK_END else float(pp))
          for d, gg, pp in zip(mf["days"], mf["glass"], mf["probe"])}
    ok(P.price(*P.counts(ov))["premium"] == gold["premium"], "SF10 the overlap's converted glass gives the same premium")
    g = np.array([x[2] for x in st["pairs"]])
    pr = np.array([x[3] for x in st["pairs"]])
    dd = pr - g
    est = {}
    bins = collections.defaultdict(list)
    for a_, b_ in zip(g, dd):
        bins[a_].append(b_)
    est["median at each reading"] = P.converter({k: float(np.median(v)) for k, v in bins.items()})
    loc = {}
    for x in sorted(set(g)):
        m = np.abs(g - x) <= 1.5
        if m.sum() >= 10 and len(set(g[m])) > 1:
            b1, a1 = np.polyfit(g[m], dd[m], 1)
            loc[x] = a1 + b1 * x
    est["local line within 1.5 C"] = P.converter(loc)
    gs, ps = np.sort(g), np.sort(pr)
    qm = {}
    for x in sorted(set(g)):
        q = (np.searchsorted(gs, x, "left") + np.searchsorted(gs, x, "right")) / 2 / len(gs)
        qm[x] = float(np.quantile(ps, q)) - x
    est["quantile matching"] = P.converter(qm)
    for lo in (25, 28, 30):
        m = g >= lo
        b1, a1 = np.polyfit(g[m], dd[m], 1)
        est["line over glass %d C and up" % lo] = (lambda a, b: lambda x: x + a + b * x)(a1, b1)
    for deg in (2, 3):
        c = np.polyfit(g, dd, deg)
        est["polynomial of degree %d" % deg] = (lambda c: lambda x: x + float(np.polyval(c, x)))(c)
    for k, f in est.items():
        h_, s_, r = P.rung(mf, f)
        ok(h_ == mc["heat"] and s_ == mc["severe"], "SF12 %s counts every season as the golden does" % k)
        ok(r["premium"] == gold["premium"] and "%.1f" % r["expected_heat"] == fig["expected_heat_days_2027_28"] and
           "%.1f" % r["expected_severe"] == fig["expected_severe_days_2027_28"], "SF12 %s files the golden" % k)
    ok(len([d for d in prov]) > 0 and all(d[5:7] in ("07", "08") for d in prov), "SF13 provisional rows July and August")
    ok(abs(gold["expected_heat"] - np.mean([gold["projected"][s][0] for s in P.SEASONS])) < 1e-9 and
       abs(gold["expected_severe"] - np.mean([gold["projected"][s][1] for s in P.SEASONS])) < 1e-9,
       "SF14 the lines' values equal the average projected counts")
    ok("filed against the day the 24 hours began" in notes and "24 hours from 9 am" in history, "SF15 the 9 am day")
    conv = P.converter(pooled)
    vals = [conv(float(x)) for d, x in zip(mf["days"], mf["glass"]) if d < W.CHANGE]
    ok(min(min(abs(v - W.HEAT_C), abs(v - W.SEVERE_C)) for v in vals) >= 0.1,
       "SF1 no converted reading within 0.1 C of a trigger")

    # ---- rounding bins
    frac = gold["premium_exact"] % 1
    ok(abs(frac - 0.5) >= 0.05, "the premium clear of a half dollar (%.3f)" % frac)
    for k in ("expected_heat", "expected_severe"):
        v = gold[k] * 10 % 1
        ok(abs(v - 0.5) >= 0.05, "%s clear of its rounding edge" % k)
    for sd, v in fig["sub_districts"].items():
        x = gold["premium"] * v["ha"]
        ok(abs(x % 1 - 0.5) > 1e-6 or "50 cents up" in terms27, "the sub-district premium's half dollar is filed (%s)" % sd)
    ok("50 cents up" in terms27, "the term sheet files the half dollar")
    ok("Its sub-districts are the ones the co-operative reports by" in _ws(notes), "Ask A the register's sub-districts govern")
    paid = collections.defaultdict(float)
    for p in st["ledger"]:
        paid[p["claim_no"]] += p["amount"]
    seas = {r["claim_no"]: r["season"] for r in st["claims"]}
    tot = collections.defaultdict(float)
    for c, v in paid.items():
        tot[seas[c]] += v
    for lab, v in tot.items():
        ok(abs(round(v, 2) % 1 - 0.5) >= 0.01, "season %s paid clear of a half dollar" % lab)

    # ---- the asks: every stop misses at every sub-district and season
    gA = {sd: (v["blocks"], v["ha"]) for sd, v in fig["sub_districts"].items()}
    for k, v in ask_a_stops(st["blocks"]).items():
        for sd in W.SUB_DISTRICTS:
            ok(v[sd] != gA[sd], "Ask A %s misses at %s" % (k, sd))
    for k in ("schedule labels, one row a block",):
        v = ask_a_stops(st["blocks"])[k]
        ok(sum(x[0] for x in v.values()) == fig["blocks_total"], "Ask A the stale labels keep the total blocks")
    gB = {lab: (v["claims"], v["amount"]) for lab, v in fig["seasons"].items()}
    for k, v in ask_b_stops(st).items():
        for lab in gB:
            ok(v[lab][0] != gB[lab][0] or v[lab][1] != gB[lab][1], "Ask B %s misses at %s" % (k, lab))
    ok(all(v["claims"] > 0 for v in fig["seasons"].values()), "every season of the cover paid")
    late_final = sum(1 for s in range(W.COVER_FIRST_SEASON, W.LAST_SEASON + 1) if W.APPROVAL[s].month >= 6)
    ok(late_final >= 6, "the final run falls in June or later for most seasons")
    claims_by = collections.defaultdict(set)
    for r in st["claims"]:
        claims_by[(r["season"], r["grower_ref"])].add(r["claim_no"])
    dups = collections.Counter(k[0] for k, v in claims_by.items() if len(v) > 1)
    ok(all(dups[lab] >= 2 for lab in gB), "two or more duplicate lodgements every season")
    singles = {c for v in claims_by.values() if len(v) == 1 for c in v}
    ok(all(c in paid for c in singles), "every single lodgement is paid")
    ok(all(sum(c in paid for c in v) == 1 for v in claims_by.values()), "each member's season paid on one claim")
    # the stale sub-district on the schedule is what the policy system stores
    sch = list(csv.DictReader(open(os.path.join(target, R.SCHEDULE), encoding="utf-8")))
    stale = sum(1 for b in BK.schedule_2027(st["blocks"]) if BK.stored_district(b) != b.district_new)
    ok(stale >= 8, "eight or more scheduled blocks carry the pre-2025 sub-district (%d)" % stale)
    ok(len(sch) > len({(r["grower_ref"], r["block_no"]) for r in sch}), "the schedule repeats blocks on variety lines")

    # ---- notes register every shipped file but themselves (H9); dates (H16)
    ok({f for f in files if f in notes} == set(files) - {D.NOTES}, "the extract notes name every shipped file")
    late = []
    for f in files:
        if f.endswith((".csv", ".md", ".json")) and f != R.SCHEDULE:
            t = open(os.path.join(target, f), encoding="utf-8").read()
            if any(x > W.AS_OF.isoformat() for x in set(re.findall(r"\b20\d\d-\d\d-\d\d\b", t))):
                late.append(f)
    ok(late == [], "no record runs past the extract date but the schedule's forward dates (%s)" % late)
    led = list(csv.DictReader(open(os.path.join(target, R.LEDGER), encoding="utf-8")))
    ok(max(dt.datetime.strptime(r["posting_date"], "%d/%m/%Y").date() for r in led) <= W.AS_OF, "ledger before extract")
    sch_late = [r for r in sch if r["inception"] > W.AS_OF.isoformat()]
    ok(all(r["inception"] == "2027-11-01" for r in sch_late), "the schedule's forward inceptions are the renewal's")

    # ---- realism sweep
    tmax = [float(r["value"]) for r in daily if r["element"] == "TMAX"]
    ok(max(tmax) < 48.5 and min(tmax) > -2.0, "the probe's range is a valley station's")
    ok(max(reg_days.values()) < 47.0, "the glass register's range is a valley station's")
    ok(len(set(tmax)) > 300, "the daily maxima vary")
    for k, (who, role) in W.PEOPLE.items():
        if k != "analyst":
            ok(any(who in t for t in texts.values()), "persona %s in the pack" % who)
    wb = openpyxl.load_workbook(os.path.join(target, R.BLOCKS), read_only=True)
    ok(wb.sheetnames == ["Blocks"], "block register sheet named")
    wr = openpyxl.load_workbook(os.path.join(target, R.REGISTER), read_only=True)
    ok(wr.sheetnames[0] == "1964" and wr.sheetnames[-1] == "1995", "a sheet a year 1964 to 1995")
    heats = [mc["heat"][s] for s in P.SEASONS]
    ok(len(set(heats)) > 12 and min(heats) >= 1, "the seasons' heat days vary")

    # ---- goldens agree with figures.json
    from pypdf import PdfReader
    qp = os.path.join(golden, "heat_cover_quote_2027_28.pdf")
    ok(len(PdfReader(qp).pages) == 2, "a two-page quote")
    q = _ws(_pdf_text(qp))
    for tok in ["$%d a hectare" % fig["call_premium_per_ha"], "%s heat days" % fig["expected_heat_days_2027_28"],
                "%s severe heat days" % fig["expected_severe_days_2027_28"], fig["conversion_at_34_5"],
                fig["conversion_at_39_0"], fig["mean_conversion_merrin_flat"], "$%d a hectare" % fig["default_premium_per_ha"],
                fig["heat_line_slope_days_per_decade"], fig["default_heat_line_slope_days_per_decade"]]:
        ok(tok in q, "quote carries %r" % tok)
    for sd, v in fig["sub_districts"].items():
        ok("%s %d %.1f ${:,}".format(v["premium"]) % (sd, v["blocks"], v["ha"]) in q, "quote row " + sd)
    for lab, v in fig["seasons"].items():
        ok("%s %d ${:,}".format(v["amount"]) % (lab, v["claims"]) in q, "quote season row " + lab)
    ok("\u2014" not in q and "\u2013" not in q, "no dash in the quote")
    rows = list(csv.DictReader(open(os.path.join(golden, "heat_cover_payouts.csv"), encoding="utf-8")))
    ok([r["season"] for r in rows] == list(fig["seasons"]), "payouts one row a season in order")
    for r in rows:
        v = fig["seasons"][r["season"]]
        ok(int(r["claims_paid"]) == v["claims"] and int(r["paid_dollars"]) == v["amount"], "payouts " + r["season"])
    print("checks: %d assertions passed (premium $%d, %s heat days, %s severe; default $%d; R2 linear $%d; R3 $%d and $%d; "
          "uniform offsets %s)" % (N[0], gold["premium"], fig["expected_heat_days_2027_28"],
                                   fig["expected_severe_days_2027_28"], rungs["R1 mean difference"]["premium"],
                                   rungs["R2 linear fit over all days"]["premium"],
                                   rungs["R3 one offset, glass at 33 C or more"]["premium"],
                                   rungs["R3 one offset, glass at 35 C or more"]["premium"],
                                   sorted(set(sweep.values()))))
