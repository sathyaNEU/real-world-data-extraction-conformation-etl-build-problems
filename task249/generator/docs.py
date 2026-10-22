"""task249 documents: the examinations manager's request (eml); the examination regulations and the item review
panel's 2027 minutes (pdf); the item bank snapshot, the invigilators' incident reports and the test centres'
capacity (xlsx); the data notes (md)."""
import datetime as dt
import os

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab import rl_config
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import records as RC


_TTF = "/usr/share/fonts/truetype/liberation"
pdfmetrics.registerFont(TTFont("LiberationSans", os.path.join(_TTF, "LiberationSans-Regular.ttf")))
pdfmetrics.registerFont(TTFont("LiberationSans-Bold", os.path.join(_TTF, "LiberationSans-Bold.ttf")))
SANS, SANS_BOLD = "LiberationSans", "LiberationSans-Bold"
rl_config.canvas_basefontname = SANS            # pages start in the document's own font


def _retag(path, tags):
    """the embedded subsets' six-letter tags, one per font, in place of the writer's AAAAAA"""
    b = open(path, "rb").read()
    for name, tag in tags:
        b = b.replace(b"AAAAAA+" + name, tag + b"+" + name)
    assert b"AAAAAA+" not in b
    open(path, "wb").write(b)


def _addr(who):
    return "%s <%s@%s>" % (who[0], who[1], P.DOMAIN)


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName=SANS, fontSize=9.6, leading=13.4, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName=SANS_BOLD, fontSize=13.5, leading=17, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName=SANS_BOLD, fontSize=10.5, leading=14,
                        spaceBefore=8, spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10.5, textColor=colors.HexColor("#444444"))
    cell = ParagraphStyle("cell", fontName=SANS, fontSize=8.0, leading=9.8)
    return body, h1, h2, small, cell


def _grid(rows, widths, cell):
    t = Table([[Paragraph(str(c), cell) for c in r] for r in rows], colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA0AB")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E3E8EF"))]))
    return t


def _footer(text):
    def foot(canvas, d):
        canvas.saveState()
        canvas.setFont(SANS, 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(20 * mm, 10 * mm, text)
        canvas.drawRightString(A4[0] - 20 * mm, 10 * mm, "Page %d" % d.page)
        canvas.restoreState()
    return foot


def _doc(path, title, author, subject):
    return SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                             bottomMargin=18 * mm, title=title, author=author, creator=author, subject=subject,
                             invariant=1)


def _sections(s, sec, body, h2):
    for head, paras in sec:
        s.append(Paragraph(head, h2))
        for p in paras:
            s.append(Paragraph(p, body))


# ------------------------------------------------------------------ the request (eml)
def request(path, F):
    text = f"""From: {_addr(P.MANAGER)}
To: {_addr(P.PSYCH)}
Date: Mon, 17 Jan 2028 09:42:17 +0530
Subject: 2027 integrity report: resit exposure
Message-ID: <20280117094217.6621@{P.DOMAIN}>
MIME-Version: 1.0
Content-Type: text/plain; charset="utf-8"
Content-Transfer-Encoding: 8bit

Abhiram,

The Examination Committee meets on {P.COMMITTEE_DATE} and the integrity report goes to
the accreditation body by {P.REPORT_DUE}. I need the resit exposure section
(regulations section 7) for the ten 2027 forms: which forms are flagged under 7.3,
so the Committee can order their reassembly, and the report's two other items
under 7.4, each form's sittings interrupted by a workstation or power failure
and those that ran in part without the centre's link.

Everything is in the shared folder: Pariksha's results exports (release 3 up to
the May window, release 4 from August), the form layouts, the bank snapshot at
the year end with the January renumbering, the register's number changes, the
review bookings, Pariksha's seat plan and event log for 2027, the centres'
incident reports, and the panel's decision log and minutes for the bank section.

Thanks,
Yahvi

--
{P.MANAGER[0]}
Examinations Manager, {P.ORG}
Chennai
"""
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(text)


# ------------------------------------------------------------------ the regulations (pdf)
def regulations(path, F):
    body, h1, h2, small, cell = _styles()
    doc = _doc(path, "%s examination: regulations" % P.EXAM, P.ORG, "Regulations")
    s = [Paragraph(P.ORG + ", Chennai", small),
         Paragraph("%s examination: regulations" % P.EXAM, h1),
         Paragraph("Approved by the Council on 12 December 2025 for the examination's first windows in 2026. "
                   "Section 3.6 amended 9 June 2027.", small), Spacer(1, 4)]
    sec = [
        ("1. The examination",
         ["The %s examination certifies that a coder can assign diagnosis and procedure codes and their modifiers "
          "from clinical documentation to the standard the Institute's members work to. It is delivered on the "
          "Institute's behalf by its test delivery provider, %s, at the Institute's six test centres (Chennai, "
          "Bengaluru, Hyderabad, Pune, Kochi and Noida)." % (P.EXAM, P.VENDOR)]),
        ("2. Registration",
         ["2.1 Each candidate holds an Institute candidate number. The number changes when the candidate's "
          "membership class changes; the register records the earlier number against the new one.",
          "2.2 A candidate books each sitting with the provider, who issues a booking reference for it."]),
        ("3. Windows, forms and sittings",
         ["3.1 The examination is offered in four windows a year, in February, May, August and November, each of "
          "twelve days with a morning and an afternoon session at every centre.",
          "3.2 Each sitting presents 170 items in a fixed order: 150 scored items and 20 trial items. Trial items are "
          "being tried for future forms and do not count towards the result. The positions of the trial items are "
          "the same on every form; the trial items themselves change from window to window.",
          "3.3 A sitting lasts four hours of examination time. Candidates may use the current ICD-10-CM, ICD-10-PCS, "
          "CPT and HCPCS code books. An item not answered when the four hours end is recorded blank and scored "
          "incorrect.",
          "3.4 A form is the set of 150 scored items in their order. Forms are assembled from the Institute's item "
          "bank. Each bank item is written from an item model, and no form carries two items written from the same "
          "item model.",
          "3.5 The delivery system draws each candidate's form at random from the forms in use in the window, "
          "leaving out any form the candidate has sat before.",
          "3.6 If a candidate's workstation (its computer, screen, keyboard or mouse) or its power fails during a "
          "sitting, the examination clock stops, and the candidate continues from their saved responses when their "
          "workstation or a spare one is ready. An invigilator may pause a sitting (for a candidate taken ill, an "
          "identity check at the desk, a change of seat or an alarm); the clock stops while the sitting is paused."]),
        ("4. Results, reviews and resits",
         ["4.1 A candidate passes with 105 or more of the 150 scored items correct.",
          "4.2 Results are released about three weeks after the window closes. A candidate who fails may book one "
          "results review within eight weeks of the result. In a supervised session the candidate is shown each "
          "scored item they answered wrongly at that sitting, with its keyed answer; items left blank are not "
          "shown. Notes may not be taken.",
          "4.3 A candidate who fails may resit in any later window."]),
        ("5. Conduct",
         ["5.1 A candidate may not copy, record or remove any item, or disclose an item's content to anyone, during "
          "or after a sitting or a results review.",
          "5.2 The Institute may withhold a result while a report of misconduct at the sitting is investigated."]),
        ("6. Trial items",
         ["6.1 Each window's trial items are new items, or revised ones, being tried for the operational bank.",
          "6.2 The item review panel meets after each window and decides on that window's trial items from their "
          "statistics among the window's first-time candidates who answered them: accept into the operational bank, "
          "return to the writer for revision, or reject. The panel may later withdraw an item it has accepted."]),
        ("7. Resit exposure",
         ["7.1 A resitting candidate's result can be raised by what an earlier sitting or results review showed them.",
          "7.2 For each form sat in a year, the Examination Committee estimates the form's exposure gain for that "
          "year: the average, over the year's resit sittings of the form (sittings by a candidate who had sat the "
          "examination before), of the number of the sitting's scored items the candidate answered correctly that "
          "they would not have answered correctly had no earlier sitting or results review shown them any item.",
          "7.3 A form whose exposure gain for the year exceeds 1.0 scored item is flagged: its resit results for that "
          "year carry the exposure flag in the integrity report, and the form is reassembled before it is used again.",
          "7.4 The integrity report also states, for each form, the year's sittings of the form that a failure of the "
          "candidate's workstation or its power interrupted, and those that ran in part without the centre's link to "
          "the provider."]),
    ]
    _sections(s, sec, body, h2)
    doc.build(s, onFirstPage=_footer("%s examination regulations" % P.EXAM),
              onLaterPages=_footer("%s examination regulations" % P.EXAM))
    _retag(path, [(b"LiberationSans-Bold", b"QHVKWB"), (b"LiberationSans", b"JRZEMT")])


# ------------------------------------------------------------------ the panel's minutes (pdf)
def minutes(path, F, counts):
    body, h1, h2, small, cell = _styles()
    doc = _doc(path, "Item review panel: minutes, 2027", P.ORG, "Minutes")
    s = [Paragraph(P.ORG + ", examinations department", small),
         Paragraph("Item review panel: minutes of the 2027 meetings (extract)", h1),
         Paragraph("Chair: the Head of Examinations. Members: three senior coders from member hospitals, the "
                   "examinations psychometrician, the bank officer (secretary).", small), Spacer(1, 4)]
    names = {4: "February", 5: "May", 6: "August", 7: "November"}
    for w in (4, 5, 6, 7):
        d = RC.PANEL[w]
        paras = ["Held at the Institute, Chennai. The panel reviewed the %d trial items tried in the %s 2027 window, "
                 "with each item's statistics among the window's first-time candidates who answered it and the "
                 "writers' notes, and decided on each item. The decisions are recorded in the bank." % (counts[w], names[w])]
        if w == 5:
            paras.append("The panel withdrew two items it accepted in April: the writers' re-check against the coding "
                         "guidance found a second defensible answer to each. Both go back to their writers.")
        if w == 6:
            paras.append("The panel withdrew one item accepted in July on the same ground.")
        if w == 7:
            paras.append("The panel confirmed its meeting dates for 2028: 11 April, 11 July, 10 October and "
                         "12 December.")
        s.append(Paragraph("%s" % d.strftime("%d %B %Y").lstrip("0"), h2))
        for p in paras:
            s.append(Paragraph(p, body))
    doc.build(s, onFirstPage=_footer("Item review panel minutes 2027"), onLaterPages=_footer("Item review panel minutes 2027"))
    _retag(path, [(b"LiberationSans-Bold", b"GXPLSA"), (b"LiberationSans", b"WNCYDF")])


# ------------------------------------------------------------------ workbooks
def _sheet(ws, header, rows, widths):
    ws.append(header)
    for c in range(1, len(header) + 1):
        ws.cell(row=1, column=c).font = Font(bold=True)
    for r in rows:
        ws.append(list(r))
    for k, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(k)].width = wdt
    ws.freeze_panes = "A2"


def bank_xlsx(path, bank):
    wb = Workbook()
    ws = wb.active
    ws.title = "Items"
    _sheet(ws, list(bank.columns), bank.itertuples(index=False), [13, 11, 11, 34, 13, 20])
    wb.save(path)


def incidents_xlsx(path, rep):
    """each centre keeps its own sheet, in its own columns, dates and names for the two sessions"""
    widths = {"Chennai": [11, 9, 7, 12, 110], "Bengaluru": [12, 11, 10, 9, 104], "Hyderabad": [11, 9, 7, 8, 100],
              "Pune": [11, 9, 7, 10, 96], "Kochi": [11, 9, 7, 12, 92], "Noida": [10, 11, 11, 7, 100]}
    wb = Workbook()
    first = True
    for c in P.CENTRES:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = c
        fmt = P.SHEETS[c]
        part = rep[rep["centre"] == c]
        rows = [(r.date.strftime(fmt["date"]), fmt["sess"][0 if r.session == "AM" else 1], r.time, r.workstation,
                 r.incident) for r in part.itertuples()]
        _sheet(ws, fmt["cols"], rows, widths[c])
    wb.save(path)


def capacity_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Centres 2027"
    rows = []
    for c in P.CENTRES:
        n = RC.WORKSTATIONS[c]
        rows.append((c, P.CENTRE_CODE[c], 1 if n <= 42 else 2, n, RC.SPARES[c], 2, max(2, n // 15),
                     "09:30" , "14:30"))
    _sheet(ws, ["Centre", "Code", "Rooms", "Workstations", "Spare workstations", "Sessions a day", "Invigilators a session",
                "Morning start", "Afternoon start"], rows, [12, 7, 7, 13, 17, 14, 21, 14, 16])
    wb.save(path)


# ------------------------------------------------------------------ data notes (md)
def notes(path, F):
    text = f"""# Examination data: notes for the 2027 review

Kept by the examinations department. The provider's files are as {P.VENDOR} delivers them.

## Results exports ({F['export3']}, {F['export4']})

One row per sitting. Release 3 of the provider's export ran to the May 2027 window; release 4 from the August 2027
window.

| Release 3 | Release 4 | Meaning |
|---|---|---|
| candidate_no | cand_no | Institute candidate number at the sitting |
| booking_ref | booking_ref | the provider's booking reference |
| window | exam_window | examination window |
| sitting_date | sat_on | date of the sitting (release 4: day-month-year) |
| centre | test_centre | test centre (release 4: the centre's code) |
| form | form_code | the form the delivery system drew |
| responses | response_string | one character per delivered position, in delivery order (170): the option chosen, A to D, or `-` where no answer was recorded |
| item_seconds | latency_ms | time on each delivered position, in delivery order, separated by `|`, empty where no answer was recorded (release 3: whole seconds; release 4: milliseconds) |
| end | end_reason | how the sitting ended: submitted by the candidate, or time expired |
| score | raw_score | scored items answered correctly (of 150) |
| result | outcome | Pass or Fail (release 4: P or F) |

## Form layouts ({F['layouts']})

Each window's forms position by position: the item delivered at each of the 170 positions and whether the position
is scored or trial. Item numbers are the bank's numbering at the time: the 2026 windows' layouts use the numbers
before the bank's migration of 11 January 2027 ({F['crosswalk']} maps them to the bank's present numbers).

## Item bank snapshot ({F['bank']})

Every item in the bank at 31 December 2027: its number, the item model it was written from, its format (case: an
operative, discharge or clinic note to code; single: a stand-alone question), content area, keyed answer and status.

## Candidate number changes ({F['register']})

The register's changes of candidate number (regulation 2.1): the earlier number, the new number and the date.

## Results review bookings ({F['reviews']})

One row per review booked: the candidate number at booking, the window of the sitting reviewed, the dates booked and
of the review, and whether it was held, cancelled or not attended, or is still to come (booked).

## Seat plan and workstation event log ({F['seat']}, {F['events']})

The provider's 2027 seat plan: each booking's centre code, date, session and the workstation it was seated at. The
event log: one row per event on a centre's workstations, time in UTC, with the booking reference of the sitting it
belongs to.

| Event | Meaning |
|---|---|
| TUTORIAL_START | the candidate starts the tutorial; the examination clock is not running |
| SESSION_START | the examination clock starts |
| WS_FAIL | the workstation stopped responding |
| SESSION_RESUME | the candidate's sitting runs again after a stop; the clock restarts |
| SESSION_END | the sitting ended (submitted, or time expired) |
| NET_LOST, NET_BACK | the centre's link to the provider lost and restored; the delivery client carries on and sends the responses when the link is back |
| WS_REBOOT | a workstation restarted |

## Incident reports ({F['incidents']})

The invigilators' reports from each centre, in local time.

## Item review panel decisions ({F['panel_log']})

The bank officer's log of the panel's decisions on trial items: the item, the decision (accept, revise, reject,
withdraw) and the date the decision was logged. Entries before 11 January 2027 use the earlier item numbers.
"""
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(text)
