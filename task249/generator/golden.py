"""The golden deliverables from figures.json: resit_exposure_review_2027.xlsx (the flags, each form's gain and its two
parts, the integrity report's two other items, notes) and resit_exposure_2027.pdf (every form's gain against the
one-item line, the flagged forms marked)."""
import json
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import params as P
import postprocess

XLSX = "resit_exposure_review_2027.xlsx"
PDF = "resit_exposure_2027.pdf"
BLUE, ORANGE, INK, INK2, GRID = "#2a78d6", "#eb6834", "#0b0b0b", "#52514e", "#d9d8d4"
WINDOWS_OF = {f: ("February and May 2027" if f < "CC27-F" else "August and November 2027") for f in P.F27}


def workbook(path, fig):
    c, a = fig["call"], fig["asks"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Exposure 2027"
    bold = Font(bold=True)
    thin = Side(style="thin", color="BFBFBF")
    ws["A1"] = "Resit exposure review, 2027 forms (regulations 7.2 and 7.3)"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Flagged: %s. Their 2027 resit results carry the exposure flag; the forms are reassembled before reuse." % (
        ", ".join(c["flagged"]))
    ws["A2"].font = bold
    ws["A3"] = "Examinations department, 24 January 2028, for the Examination Committee of %s." % P.COMMITTEE_DATE
    ws["A3"].font = Font(italic=True, size=9, color="52514E")
    head = ["Form", "Windows", "Exposure gain per resit sitting (scored items)",
            "of which on items shown before", "of which in the time those items freed", "Flag"]
    ws.append(head)
    for k in range(1, len(head) + 1):
        cl = ws.cell(row=4, column=k)
        cl.font = bold
        cl.alignment = Alignment(wrap_text=True, vertical="top")
        cl.border = Border(bottom=thin)
    for f in P.F27:
        g = c["gain"][f]
        ws.append([f, WINDOWS_OF[f], round(g, 1), round(c["direct"][f], 1), round(c["reach"][f], 1),
                   "Flagged" if f in c["flagged"] else ""])
    for r in range(5, 5 + len(P.F27)):
        for k in (3, 4, 5):
            ws.cell(row=r, column=k).number_format = "0.0"
        if ws.cell(row=r, column=6).value:
            for k in range(1, 7):
                ws.cell(row=r, column=k).fill = PatternFill("solid", fgColor="EEF4FC")
    for k, wdt in enumerate([10, 24, 22, 18, 20, 10], 1):
        ws.column_dimensions[get_column_letter(k)].width = wdt
    ws.row_dimensions[4].height = 46
    ws.freeze_panes = "A5"
    ws2 = wb.create_sheet("Sittings by form")
    ws2.append(["Form", "2027 sittings a failure of the workstation or its power interrupted",
                "2027 sittings that ran in part without the centre's link"])
    for k in range(1, 4):
        ws2.cell(row=1, column=k).font = bold
        ws2.cell(row=1, column=k).alignment = Alignment(wrap_text=True, vertical="top")
    for f in P.F27:
        ws2.append([f, a["interrupted"][f], a["offline"][f]])
    ws2.append(["All forms", sum(a["interrupted"].values()), sum(a["offline"].values())])
    for k in range(1, 4):
        ws2.cell(row=ws2.max_row, column=k).font = bold
        ws2.cell(row=ws2.max_row, column=k).border = Border(top=thin)
    ws2.append([])
    ws2.append(["Event log read by workstation and time; the Kochi power cut of %s from the invigilator's report." %
                fig["asks"]["power_day"]])
    ws2.cell(row=ws2.max_row, column=1).font = Font(italic=True, size=9, color="52514E")
    for k, wdt in enumerate([10, 30, 30], 1):
        ws2.column_dimensions[get_column_letter(k)].width = wdt
    ws2.row_dimensions[1].height = 32
    ws2.freeze_panes = "A2"
    ws4 = wb.create_sheet("Notes")
    notes = [
        "Exposure gain (regulation 7.2): for each resit sitting in 2027, the scored items answered correctly that would "
        "not have been had no earlier sitting or review shown the candidate any item, averaged over the form's 2027 "
        "resit sittings. A resit sitting is a sitting by a candidate who had sat before under any of their candidate "
        "numbers (register of number changes).",
        "Items shown before: the same item reached at an earlier sitting (scored or trial position), or an item written "
        "from the same item model; held reviews show again items the candidate had reached and answered wrongly.",
        "On items shown before: the candidate's correct answers on those items less what their ability on the other "
        "scored items gives, against the window's first-time candidates' share correct on each item.",
        "In the time those items freed: an item recognised from an earlier sitting takes less time than the candidate's "
        "own pace on the other items gives (the first-time candidates' mean log time on the item, shifted by the "
        "candidate's); the correct answers on the items reached only within that time count as gain.",
        "Interrupted sittings (7.4): a failure of the candidate's workstation, from the event log read by workstation and "
        "time, and the Kochi power cut of %s from the invigilator's report; invigilators' pauses and network losses "
        "are not failures." % fig["asks"]["power_day"],
        "Sittings that ran in part without the link (7.4): the examination clock running while the workstation had lost "
        "the centre's link, from each NET_LOST to the next NET_BACK on that workstation, or to the sitting's end where the "
        "link did not come back while it ran.",
        "Sources: Pariksha results exports releases 3 and 4, form layouts, item bank snapshot of 31 December 2027 and the "
        "January renumbering, candidate number changes, review bookings, seat plan and workstation event log, "
        "invigilator incident reports.",
    ]
    ws4.column_dimensions["A"].width = 120
    for t in notes:
        ws4.append([t])
        ws4.cell(row=ws4.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def chart(path, fig):
    c = fig["call"]
    forms = P.F27[::-1]
    d = [c["direct"][f] for f in forms]
    r = [c["reach"][f] for f in forms]
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9})
    figr, ax = plt.subplots(figsize=(8.3, 5.4))
    figr.patch.set_facecolor("#fcfcfb")
    ax.set_facecolor("#fcfcfb")
    y = range(len(forms))
    ax.barh(y, d, height=0.56, color=BLUE, label="on items shown before", edgecolor="#fcfcfb", linewidth=1.5)
    ax.barh(y, r, left=d, height=0.56, color=ORANGE, label="in the time those items freed", edgecolor="#fcfcfb",
            linewidth=1.5)
    ax.axvline(P.THRESH, color=INK, linewidth=1.0, linestyle=(0, (4, 3)))
    ax.set_ylim(-0.95, len(forms) - 0.45)
    ax.text(P.THRESH + 0.03, -0.72, "1.0 item, the line in regulation 7.3", color=INK, fontsize=8, va="center")
    for k, f in enumerate(forms):
        g = c["gain"][f]
        tag = "  flagged" if f in c["flagged"] else ""
        ax.text(g + 0.04, k, "%.1f%s" % (g, tag), va="center", fontsize=8.5, color=INK if tag else INK2,
                fontweight="bold" if tag else "normal")
    ax.set_yticks(list(y))
    ax.set_yticklabels(forms)
    for lab, f in zip(ax.get_yticklabels(), forms):
        if f in c["flagged"]:
            lab.set_fontweight("bold")
    ax.set_xlabel("Exposure gain per 2027 resit sitting (scored items)", color=INK2)
    ax.set_xlim(0, max(c["gain"].values()) + 0.75)
    ax.grid(axis="x", color=GRID, linewidth=0.6)
    ax.set_axisbelow(True)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.legend(loc="lower right", frameon=False, fontsize=8)
    figr.suptitle("Six of the ten 2027 forms gave their resit candidates more than one scored item", x=0.06, ha="left",
                  fontsize=11.5, fontweight="bold", color=INK, y=0.97)
    ax.set_title("Each form's gain per resit sitting: on the items an earlier sitting or review had shown, and in the "
                 "time those items freed", loc="left", fontsize=8.5, color=INK2)
    figr.text(0.06, 0.015, "IHCP examinations department, integrity report 2027. Sources: Pariksha results exports, form "
              "layouts, item bank snapshot of 31 December 2027.", fontsize=7, color=INK2)
    figr.tight_layout(rect=(0, 0.03, 1, 0.94))
    with PdfPages(path, metadata={"Title": "Resit exposure 2027", "Author": P.PSYCH[0], "Creator": P.ORG,
                                  "Producer": P.ORG, "CreationDate": None, "ModDate": None}) as pdf:
        pdf.savefig(figr, facecolor=figr.get_facecolor())
    plt.close(figr)


def write(outdir, fig):
    os.makedirs(outdir, exist_ok=True)
    x = os.path.join(outdir, XLSX)
    workbook(x, fig)
    postprocess.scrub(x, P.ORG, "2028-01-24")
    postprocess.finish_office(x, "2028-01-19T05:10:00Z", "2028-01-24T11:35:00Z", P.ORG)
    p = os.path.join(outdir, PDF)
    chart(p, fig)
    postprocess.scrub(p, P.ORG, "2028-01-24T11:40:00")
