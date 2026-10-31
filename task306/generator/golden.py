"""The two golden deliverables: the actuarial function's workbook and the board paper."""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import postprocess
import world as W

XLSX = "windstorm_programme_2027.xlsx"
PDF = "windstorm_programme_board_paper_2027.pdf"


def m1(x):
    """pounds to whole millions, the precision the capital figures are given at"""
    return int(round(x / 1e6))


def m01(x):
    return round(x / 1e6, 1)


def k1(x):
    """pounds to the nearest thousand"""
    return int(round(x / 1000.0)) * 1000


def occ_label(o, year=True):
    s, e = o["start"], o["end"]
    if s == e:
        return s.strftime("%d %b %Y") if year else s.strftime("%d %b")
    if s.month == e.month:
        return "%d to %s" % (s.day, e.strftime("%d %b %Y" if year else "%d %b"))
    return "%s to %s" % (s.strftime("%d %b"), e.strftime("%d %b %Y" if year else "%d %b"))


def workbook(path, G):
    wb = Workbook()
    ws = wb.active
    ws.title = "Capital 2027"
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="E3EAEE")
    ws.append(["Windstorm capital under CP-04, 2027 portfolio"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Seaborne UK Windstorm Model 8.2, 10,000 years; 1-in-200 = average of the 50 worst years"])
    ws.append([])
    ws.append(["", "GBP m", "Note"])
    for c in ws[4]:
        c.font = bold
        c.fill = fill
    rows = [
        ("1-in-200 annual windstorm loss, gross", m1(G["gross"]), "", "0"),
        ("1-in-200, net of the occurrence cover", m1(G["net_base"]), "18m xs 12m, one reinstatement", "0"),
        ("1-in-200, net of the occurrence cover and the second-event cover", m1(G["net_both"]),
         "8m xs 4m of the retained loss, second and later occurrences", "0"),
        ("Capital released by the second-event cover", m1(G["release_sec"]), "", "0"),
        None,
        ("Second-event cover: premium", m01(W.SEC["premium"]), "Penwarden quotation, cover B", "0.0"),
        ("  expected reinstatement premium", m01(G["rp_sec"]), "average over the 10,000 years", "0.0"),
        ("  expected recoveries", m01(G["rec_sec"]), "average over the 10,000 years", "0.0"),
        ("  margin", m01(G["margin_sec"]), "premium and reinstatement premium less recoveries", "0.0"),
        ("Cost of the capital released at 12 per cent", m01(W.COC * G["release_sec"]), "", "0.0"),
        ("Value of the second-event cover", m01(G["value"]), "12 per cent of the release less the margin", "0.0"),
    ]
    for r in rows:
        if r is None:
            ws.append([])
            continue
        ws.append(list(r[:3]))
        ws.cell(row=ws.max_row, column=2).number_format = r[3]
    ws.append([])
    ws.append(["Recommendation", "Add the second-event cover to the 2027 programme"])
    ws["A%d" % ws.max_row].font = bold
    ws.append([])
    ws.append(["Each table storm costed at incurred over hindcast for the storms of record like it"])
    ws["A%d" % ws.max_row].font = bold
    nf = sum(G["record_follow"])
    for label, v, note in (("  storms within six weeks of the storm before", G["load_follow"], "%d storms of record" % nf),
                           ("  every other storm", G["load_lone"], "%d storms of record" % (len(G["record"]) - nf)),
                           ("  all storms of record together, for comparison", G["pooled"], "not used")):
        ws.append([label, round(v, 3), note])
        ws.cell(row=ws.max_row, column=2).number_format = "0.000"
    for col, w in zip("ABC", (64, 14, 58)):
        ws.column_dimensions[col].width = w

    s2 = wb.create_sheet("Storms 2026")
    s2.append(["Loss occurrences of GBP 1m or more under the 2026 occurrence cover, claims extract of 31 October 2026"])
    s2["A1"].font = bold
    s2.append([])
    s2.append(["Days of damage", "Home GBP", "Farm GBP", "Incurred GBP", "Recovery GBP", "Reinstatement premium GBP"])
    for c in s2[3]:
        c.font = bold
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for o in G["occ2026"]:
        s2.append([occ_label(o), k1(o["home"]), k1(o["farm"]), k1(o["incurred"]), k1(o["recovery"]),
                   k1(o["reinstatement"])])
    s2.append(["2026", None, None, None, k1(G["recovery2026"]), k1(G["reinstatement2026"])])
    s2["A%d" % s2.max_row].font = bold
    for k in range(4, s2.max_row + 1):
        for j in range(2, 7):
            s2.cell(row=k, column=j).number_format = "#,##0"
    s2.append([])
    s2.append(["Claims dated by the adjusters' date of damage; a reopened claim with a second instruction is dated by it."])
    for col, w in zip("ABCDEF", (24, 14, 14, 15, 15, 17)):
        s2.column_dimensions[col].width = w

    s3 = wb.create_sheet("Farm sums insured")
    s3.append(["Farm book, our own share, at 31 October 2026, by the county each holding is in"])
    s3["A1"].font = bold
    s3.append([])
    s3.append(["County", "Buildings GBP m", "Farm contents GBP m"])
    for c in s3[3]:
        c.font = bold
        c.fill = fill
    for county in W.COUNTIES:
        b, c = G["farm_si"][county]
        s3.append([county, m01(b), m01(c)])
    for k in range(4, s3.max_row + 1):
        for j in (2, 3):
            s3.cell(row=k, column=j).number_format = "#,##0.0"
    for col, w in zip("ABC", (14, 17, 20)):
        s3.column_dimensions[col].width = w

    s4 = wb.create_sheet("Record")
    s4.append(["First damage date", "Storm", "Days since the storm before", "Hindcast GBP", "Incurred GBP",
               "Incurred over hindcast", "Within six weeks of the storm before"])
    for c in s4[1]:
        c.font = bold
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
    prev = None
    for (d, h, a, name), f in zip(G["record"], G["record_follow"]):
        s4.append([d, name or None, (d - prev).days if prev else None, round(h), round(a), round(a / h, 3),
                   "Yes" if f else "No"])
        prev = d
    for k in range(2, s4.max_row + 1):
        s4.cell(row=k, column=1).number_format = "dd/mm/yyyy"
        s4.cell(row=k, column=4).number_format = "#,##0"
        s4.cell(row=k, column=5).number_format = "#,##0"
        s4.cell(row=k, column=6).number_format = "0.000"
    for col, w in zip("ABCDEFG", (14, 11, 13, 14, 14, 12, 16)):
        s4.column_dimensions[col].width = w
    s4.freeze_panes = "A2"
    wb.properties.creator = W.PEOPLE["actuary"]
    wb.properties.lastModifiedBy = W.PEOPLE["actuary"]
    wb.save(path)
    postprocess.scrub(path, W.ORG, "2026-11-04")
    postprocess.finish_office(path, "2026-11-02T08:20:00Z", "2026-11-04T16:48:00Z", W.ORG)


def chart(G):
    d = Drawing(160 * mm, 66 * mm)
    bc = VerticalBarChart()
    bc.x, bc.y, bc.width, bc.height = 16 * mm, 10 * mm, 136 * mm, 44 * mm
    vals = [G["gross"] / 1e6, G["net_base"] / 1e6, G["net_both"] / 1e6]
    bc.data = [vals]
    bc.categoryAxis.categoryNames = ["Gross", "Net of the occurrence cover", "Net of both covers"]
    bc.categoryAxis.labels.fontSize = 8
    bc.valueAxis.valueMin = 0
    bc.valueAxis.valueMax = 20 * (int(max(vals) / 20) + 1)
    bc.valueAxis.valueStep = 20
    bc.valueAxis.labels.fontSize = 8
    bc.bars[0].fillColor = colors.HexColor("#2f5d7c")
    bc.barWidth = 10
    d.add(bc)
    for i, v in enumerate(vals):
        x = bc.x + bc.width * (i + 0.5) / 3
        y = bc.y + bc.height * v / bc.valueAxis.valueMax + 2
        d.add(String(x, y, "%d" % int(round(v)), fontSize=8.5, textAnchor="middle"))
    d.add(String(16 * mm, 60 * mm, "1-in-200 annual windstorm loss, 2027, GBP million", fontSize=9))
    return d


def paper(path, G):
    st = ParagraphStyle("p", fontName="Helvetica", fontSize=9.6, leading=13.2, spaceAfter=6)
    hd = ParagraphStyle("h", parent=st, fontName="Helvetica-Bold", fontSize=10.5, spaceBefore=6)
    tt = ParagraphStyle("t", parent=st, fontName="Helvetica-Bold", fontSize=13.5, leading=17)
    sm = ParagraphStyle("s", parent=st, fontSize=8.4, leading=11, textColor=colors.HexColor("#4a5a66"))
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                            bottomMargin=18 * mm, title="2027 windstorm programme", author=W.PEOPLE["reinsurance"],
                            creator=W.ORG, producer=W.ORG, invariant=1)
    nf = sum(G["record_follow"])
    story = [
        Paragraph("Board paper: the 2027 windstorm programme", tt),
        Paragraph("For the Board meeting of November 2026. %s, Head of Reinsurance, with %s, Chief Actuary."
                  % (W.PEOPLE["reinsurance"], W.PEOPLE["actuary"]), sm),
        Paragraph("Recommendation", hd),
        Paragraph("Renew the occurrence cover and add the second-event cover Penwarden has quoted. Under CP-04 the "
                  "second-event cover takes the 1-in-200 windstorm loss net of reinsurance from GBP %d million to GBP "
                  "%d million, releasing GBP %d million of capital. At 12 per cent that capital costs GBP %.1f million "
                  "a year, against a margin of GBP %.1f million on the cover, so the cover adds GBP %.1f million a year."
                  % (m1(G["net_base"]), m1(G["net_both"]), m1(G["release_sec"]), m01(W.COC * G["release_sec"]),
                     m01(G["margin_sec"]), m01(G["value"])), st),
        Paragraph("What the cover is worth turns on what our second storms cost", hd),
        Paragraph("CP-04 costs each storm in the model's table at what our own claims show we pay for such a storm. "
                  "Our record since 2001 holds %d storms. The %d that came within six weeks of the storm before cost us "
                  "%.2f times the model's hindcast; the rest cost %.2f times. Taken together the record gives %.2f, "
                  "and one figure for every storm undercosts the storms the second-event cover pays on."
                  % (len(G["record"]), nf, G["load_follow"], G["load_lone"], G["pooled"]), st),
        chart(G),
        Spacer(1, 4),
        Paragraph("This year's storms under the expiring cover", hd),
    ]
    t = [["Days of damage", "Home GBP", "Farm GBP", "Incurred GBP", "Recovery GBP", "Reinstatement GBP"]]
    for o in G["occ2026"]:
        t.append([occ_label(o, year=False), format(k1(o["home"]), ","), format(k1(o["farm"]), ","),
                  format(k1(o["incurred"]), ","), format(k1(o["recovery"]), ","), format(k1(o["reinstatement"]), ",")])
    t.append(["2026", "", "", "", format(k1(G["recovery2026"]), ","), format(k1(G["reinstatement2026"]), ",")])
    tbl = Table(t, hAlign="LEFT")
    tbl.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 8.6), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.6),
                             ("FONT", (0, -1), (-1, -1), "Helvetica-Bold", 8.6),
                             ("ALIGN", (1, 1), (-1, -1), "RIGHT"), ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.grey),
                             ("LINEABOVE", (0, -1), (-1, -1), 0.4, colors.grey)]))
    story += [tbl, Spacer(1, 4),
              Paragraph("Loss occurrences of GBP 1 million or more, each claim dated by the day its damage was done.", sm),
              Paragraph("Farm sums insured, our own share, at 31 October 2026", hd)]
    t2 = [["County", "Buildings GBP m", "Farm contents GBP m"]] + [
        [c, "%.1f" % m01(G["farm_si"][c][0]), "%.1f" % m01(G["farm_si"][c][1])] for c in W.COUNTIES]
    tbl2 = Table(t2, hAlign="LEFT")
    tbl2.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 8.6), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.6),
                              ("ALIGN", (1, 1), (-1, -1), "RIGHT"), ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.grey)]))
    story.append(tbl2)
    doc.build(story)
    postprocess.scrub(path, W.ORG, "2026-11-04")


def write(golden, G):
    os.makedirs(golden, exist_ok=True)
    workbook(os.path.join(golden, XLSX), G)
    paper(os.path.join(golden, PDF), G)
