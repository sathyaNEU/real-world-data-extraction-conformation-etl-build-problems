"""The pack's data files, each in the layout of the system that produced it."""
import csv
import datetime as dt
import math

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import postprocess
import season as S
import world as W

FILES = {
    "ylt": "seaborne_windstorm_ylt_2027_portfolio.csv",
    "hindcast": "seaborne_windstorm_hindcast_2001-2025.csv",
    "experience": "windstorm_experience_2001-2025.xlsx",
    "capital": "CP-04_capital_policy_windstorm_2026.pdf",
    "quote": "penwarden_windstorm_quotation_2027.pdf",
    "wording": "windstorm_xl_wording_2026.pdf",
    "request": "2027_windstorm_programme.eml",
    "claims": "claims_extract_2026-10-31.csv",
    "adjusters": "loss_adjuster_instructions_2026.csv",
    "home": "home_policies_2026-10-31.csv",
    "farm": "farm_policies_and_locations_2026-10-31.xlsx",
    "coins": "coinsurance_register_2026.csv",
    "flood": "seaborne_flood_ylt_2027_portfolio.csv",
    "panel": "reinsurance_panel_security_2027.xlsx",
}
DISTRACTORS = ["flood", "panel", "home"]



def _csv(path, header, rows, sep=","):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, delimiter=sep, lineterminator="\r\n")
        w.writerow(header)
        for r in rows:
            w.writerow(r)
    return len(rows)


def _xlsx(path, sheets, made, saved, author):
    wb = Workbook()
    wb.remove(wb.active)
    for title, header, rows, widths, fmts in sheets:
        ws = wb.create_sheet(title)
        ws.append(header)
        for j, _ in enumerate(header, 1):
            c = ws.cell(row=1, column=j)
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDE5EA")
            c.alignment = Alignment(vertical="top", wrap_text=True)
        for r in rows:
            ws.append(list(r))
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        for col, f in (fmts or {}).items():
            j = header.index(col) + 1
            for k in range(2, len(rows) + 2):
                ws.cell(row=k, column=j).number_format = f
        ws.freeze_panes = "A2"
    wb.properties.creator = author
    wb.properties.lastModifiedBy = author
    wb.save(path)
    postprocess.scrub(path, W.ORG, saved[:10])
    postprocess.finish_office(path, made, saved, W.ORG)


# invented storm names, one list a season from autumn 2015, when storms began to be named: a season's names run A to W
# (no Q, U, X, Y or Z), and the mutual's storms are the few of each season's that reached it
POOL = {
    "A": ["Anwen", "Aled", "Ailsa", "Arwel", "Annis", "Aeron", "Alys", "Amory", "Aisling", "Arlo", "Avery", "Aelwen"],
    "B": ["Bryok", "Bethan", "Bodmin", "Brannoc", "Briallen", "Bram", "Branwen", "Bronte", "Beren", "Blythe", "Benno"],
    "C": ["Cador", "Carys", "Corin", "Ciaran", "Cerys", "Clovis", "Cadan", "Cressida", "Calum", "Columb", "Cari"],
    "D": ["Derwa", "Demelza", "Dilwyn", "Dorcas", "Dafydd", "Delyth", "Deverel", "Dunstan", "Dulcie", "Drew", "Dilys"],
    "E": ["Elowen", "Eseld", "Emrys", "Eira", "Esmond", "Efa", "Ellery", "Elfed", "Enid", "Ewan", "Edern"],
    "F": ["Fenwick", "Fenella", "Fergal", "Freya", "Fulke", "Flora", "Faron", "Fiona", "Fenn", "Fflur", "Faisal"],
    "G": ["Gwennol", "Garrow", "Gwilym", "Gilda", "Garnet", "Glenys", "Gowan", "Greer", "Guthrie", "Gwen", "Gideon"],
    "H": ["Hedra", "Hesketh", "Hywel", "Hester", "Huw", "Halcyon", "Harlan", "Heulwen", "Hamish", "Hafren", "Hilde"],
    "I": ["Ithel", "Ilar", "Isolde", "Idris", "Ianto", "Ione", "Islay", "Ivo", "Ilsa", "Irwin", "Ifan"],
    "J": ["Jory", "Jenifer", "Jethro", "Jocasta", "Jasper", "Jowan", "Jessamy", "Joss", "Jemima", "Jarvis", "Jac"],
    "K": ["Kenwyn", "Keverne", "Kerensa", "Kester", "Kit", "Kilda", "Kai", "Kelyn", "Kenelm", "Kezia", "Kyffin"],
    "L": ["Lowen", "Lamorna", "Llew", "Lettice", "Lorcan", "Lark", "Linnet", "Lyle", "Leofric", "Lowri", "Lazlo"],
    "M": ["Morwen", "Madron", "Meraud", "Mabyn", "Mostyn", "Marlow", "Mairwen", "Merrin", "Mungo", "Myfanwy", "Milo"],
    "N": ["Nectan", "Nonna", "Nye", "Noreen", "Niamh", "Nevin", "Nerys", "Nolan", "Nesta", "Neville", "Nell"],
    "O": ["Ottery", "Olwen", "Orrin", "Osian", "Odette", "Oswin", "Ondine", "Orla", "Otto", "Owain", "Opal"],
    "P": ["Petroc", "Piran", "Perran", "Pryce", "Peronel", "Phelim", "Posy", "Padarn", "Pippa", "Porter", "Pella"],
    "R": ["Rosen", "Ruan", "Rhodri", "Rowena", "Rhian", "Ralston", "Romilly", "Rufus", "Rhosyn", "Ravi", "Rilla"],
    "S": ["Senara", "Sennen", "Sulien", "Selwyn", "Saffron", "Stellan", "Sian", "Solly", "Seren", "Silas", "Sabra"],
    "T": ["Tamsin", "Tregony", "Talan", "Teilo", "Tamar", "Thorne", "Tegan", "Tobias", "Tirion", "Tansy", "Tybalt"],
    "V": ["Veryan", "Verity", "Vaughan", "Vashti", "Vane", "Vivian", "Velma", "Vincent", "Vala", "Vere", "Viggo"],
    "W": ["Wenna", "Wendron", "Winnoc", "Wynne", "Wren", "Wilfred", "Willa", "Wystan", "Wanda", "Wolfe", "Wini"],
}
LETTERS = "ABCDEFGHIJKLMNOPRSTVW"


def _ids(rng, n, start, spread):
    """n increasing numbers from start, in uneven steps the way a numbering system leaves them"""
    steps = rng.integers(1, spread + 1, n)
    return [int(x) for x in start + np.cumsum(steps)]


def _catalogue(rng, n, lo, hi):
    """n distinct catalogue ids, as a vendor's event set numbers its events"""
    return [int(x) for x in rng.choice(np.arange(lo, hi), size=n, replace=False)]


def _peak(rng, day, year=2027):
    base = dt.datetime(year, 1, 1) + dt.timedelta(days=day - 1)
    return base + dt.timedelta(minutes=int(rng.integers(0, 24 * 60)))


def write_ylt(path, rows):
    """the vendor's year loss table for the 2027 portfolio"""
    rng = np.random.default_rng(W.SEED + 501)
    ids = _catalogue(np.random.default_rng(W.SEED + 511), len(rows), 1_000_000, 9_999_999)
    out = []
    for k, (y, d, l, f) in enumerate(rows):
        out.append([y, "SR8-%07d" % ids[k], _peak(rng, d).strftime("%Y-%m-%dT%H:%MZ"), "%.0f" % l])
    return _csv(path, ["year", "event_id", "peak_utc", "loss_gbp"], out)


def write_flood(path):
    """the flood model's year loss table (CP-05's peril)"""
    rng = np.random.default_rng(W.SEED + 502)
    rows = []
    for y in range(1, W.N_YEARS + 1):
        for _ in range(int(rng.poisson(0.55))):
            d = int(rng.integers(1, 366))
            l = float(math.exp(rng.normal(math.log(0.6e6), 1.25)))
            if l >= W.FLOOR:
                rows.append((y, d, l))
    ids = _catalogue(np.random.default_rng(W.SEED + 512), len(rows), 2_000_000, 7_999_999)
    out = [[y, "SF3-%07d" % ids[k], _peak(rng, d).strftime("%Y-%m-%dT%H:%MZ"), "%.0f" % l]
           for k, (y, d, l) in enumerate(rows)]
    return _csv(path, ["year", "event_id", "peak_utc", "loss_gbp"], out)


def _season(d):
    return d.year if d.month >= 7 else d.year - 1


def record_rows(rec):
    """the record's storms with their names (named from the autumn 2015 season) and references"""
    rng = np.random.default_rng(W.SEED + 505)
    named, letter = [], {}
    for i, (date, hind, act, f) in enumerate(rec):
        name = ""
        if date >= dt.date(2015, 10, 1):
            s = _season(date)
            j = letter.get(s, -1) + 1 + int(rng.choice([0, 0, 1, 1, 2, 3]))
            letter[s] = j
            assert j < len(LETTERS), (date, j)
            # the 2025/26 season's names from F on are its January and February 2026 storms
            assert s != 2025 or LETTERS[j] < "F", (date, LETTERS[j])
            pool = POOL[LETTERS[j]]
            name = pool[(s - 2015) % len(pool)]
        named.append((date, hind, act, f, name, "HWS-%d-%02d" % (date.year, 1 + sum(1 for r in rec[:i] if r[0].year == date.year))))
    return named


def write_hindcast(path, named):
    rng = np.random.default_rng(W.SEED + 503)
    rid = np.random.default_rng(W.SEED + 513)
    out, last = [], {}
    for k, (date, hind, act, f, name, ref) in enumerate(named):
        # the model's peak, in UTC; a storm whose winds peaked after midnight carries the next day
        peak = dt.datetime.combine(date, dt.time(0, 0)) + dt.timedelta(minutes=int(rng.integers(6 * 60, 30 * 60)))
        # the vendor's historical event set numbers each year's events, most of them storms that missed the portfolio
        n = last.get(date.year, int(rid.integers(1, 9))) + int(rid.integers(3, 40))
        last[date.year] = n
        out.append(["SRH-%d-%03d" % (date.year, n), peak.strftime("%Y-%m-%dT%H:%MZ"), date.year, "%.0f" % hind])
    return _csv(path, ["event_id", "peak_utc", "portfolio_year", "loss_gbp"], out)


def write_experience(path, named):
    rng = np.random.default_rng(W.SEED + 504)
    rows = []
    for date, hind, act, f, name, ref in named:
        claims = int(max(12, act / float(rng.uniform(3600, 6400))))
        system = "Legacy (Sapphire)" if date < dt.date(2014, 4, 1) else "ClaimCentre"
        rows.append([ref, name or None, date, claims, round(act), system])
    _xlsx(path, [("Storms", ["storm_ref", "storm_name", "first_damage_date", "claims", "incurred_gbp",
                             "claims_system"], rows, [14, 14, 16, 9, 15, 18],
                  {"first_damage_date": "dd/mm/yyyy", "incurred_gbp": "#,##0"})],
          "2026-06-08T09:12:00Z", "2026-07-03T15:40:00Z", W.PEOPLE["actuary"])
    return len(rows)


def write_claims(path, claims):
    out = [[c["claim"], c["line"], c["peril"], c["loss_date"].strftime("%d/%m/%Y"), c["notified"].strftime("%d/%m/%Y"),
            c["channel"], c["county"], "%.2f" % c["paid"], "%.2f" % c["outstanding"], c["status"]] for c in claims]
    return _csv(path, ["claim_number", "line", "peril", "loss_date", "notified_date", "channel", "county", "paid_gbp",
                       "outstanding_gbp", "status"], out)


def write_adjusters(path, adj):
    """the loss adjusters' instruction log, in their system's layout: the insurer's claim number spaced as the
    adjusters' portal prints it, one row per instruction"""
    out = []
    for a in adj:
        c = a["claim"]
        insurer_ref = "%s %s %s" % (c[:2], c[2:6], c[6:])
        out.append([a["job"], insurer_ref, a["damage"].isoformat(), a["instructed"].isoformat(),
                    "Penhallow Loss Adjusters"])
    return _csv(path, ["job_ref", "insurer_ref", "date_of_damage", "instructed", "adjuster"], out)


def write_home(path, home):
    out = [[h["policy"], h["status"], h["start"].strftime("%d/%m/%Y"), h["end"].strftime("%d/%m/%Y"), h["county"],
            "%.0f" % h["buildings"], "%.0f" % h["contents"]] for h in sorted(home, key=lambda x: x["policy"])]
    return _csv(path, ["policy_number", "status", "start_date", "end_date", "county", "buildings_si_gbp",
                       "contents_si_gbp"], out)


def write_farm(path, farms, locations):
    pol = [[f["policy"], f["cph"], f["county"], f["si_k"]] for f in farms]
    loc = [[l["policy"], l["location"], l["cph"], l["prev"] or None, l["county"], l["b_k"], l["c_k"]] for l in locations]
    _xlsx(path, [("Policies", ["policy", "main_holding_cph", "correspondence_county", "total_si_000"], pol,
                  [12, 16, 20, 14], {"total_si_000": "#,##0"}),
                 ("Holdings", ["policy", "location_no", "holding_cph", "previous_cph", "county", "buildings_si_000",
                               "farm_contents_si_000"], loc, [12, 12, 14, 14, 12, 16, 18],
                  {"buildings_si_000": "#,##0", "farm_contents_si_000": "#,##0"})],
          "2026-10-31T16:05:00Z", "2026-10-31T16:22:00Z", W.ORG)
    return len(pol) + len(loc)


def write_coins(path, coins):
    out = [[c["lead"], c["lead_ref"], c["cph"], "%d%%" % round(c["share"] * 100), c["arranged"].strftime("%d/%m/%Y")]
           for c in coins]
    return _csv(path, ["lead_insurer", "lead_reference", "holding_cph", "hartcombe_share", "last_renewed"], out)


def write_panel(path):
    rows = [["Aldwych Syndicate 4127", "A+ (Lloyd's)", "35.0%", "35.0%"],
            ["Brunel Re Europe SE", "AA-", "20.0%", "22.5%"],
            ["Calvert Mutual Reinsurance", "A", "15.0%", "12.5%"],
            ["Dunmore Re plc", "A", "15.0%", "15.0%"],
            ["Elmstead Syndicate 1880", "A+ (Lloyd's)", "15.0%", "15.0%"]]
    _xlsx(path, [("Panel", ["reinsurer", "financial_strength", "share_2026", "share_2027_quoted"], rows,
                  [30, 18, 12, 18], {})], "2026-10-14T10:02:00Z", "2026-10-14T10:20:00Z", W.BROKER)
    return len(rows)
