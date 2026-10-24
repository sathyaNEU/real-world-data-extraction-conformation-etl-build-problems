"""task306 documents: the draft 2027 tooling plan (docx), the warranty terms summary (pdf), the RDW extract's field
descriptions (pdf), the line exports' notes (md), and the two distractors (the 2026 price list, pdf; module lead
times, xlsx)."""
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.4, leading=13.0, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=13.2, leading=17, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.2, leading=14,
                        spaceBefore=8, spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8.0, leading=10.6, textColor=colors.HexColor("#444444"))
    return body, h1, h2, small


def _footer(text):
    def foot(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.4)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(18 * mm, 10 * mm, text)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, "%d" % d.page)
        canvas.restoreState()
    return foot


def _table(rows, widths, size=8.2):
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=size, leading=size + 2.2)
    head = ParagraphStyle("head", parent=cell, fontName="Helvetica-Bold")
    rows = [[Paragraph(str(v), head if r == 0 else cell) for v in row] for r, row in enumerate(rows)]
    t = Table(rows, colWidths=[w * mm for w in widths], repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9A9A9A")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6ECEF")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


def _pdf(path, title, subject, story, foot):
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title=title, author=P.ORG, creator=P.ORG, subject=subject,
                            invariant=1)
    doc.build(story, onFirstPage=_footer(foot), onLaterPages=_footer(foot))


def warranty(path):
    body, h1, h2, small = _styles()
    F = P.FILES
    rows = [["Platform", "OEM pack warranty", "Runs from", "Mileage limit"]]
    for n in P.NAMES:
        rows.append([n, "%d years" % P.WARRANTY[n], "datum eerste toelating", "none"])
    s = [Paragraph(P.ORG + " | Technical Services", small), Spacer(1, 2 * mm),
         Paragraph("Traction-battery warranty terms on the five platforms we refurbish", h1),
         Paragraph("Summary for planning and the service desk, 3 February 2026. Replaces the summary of March 2024. "
                   "Sources: each OEM's warranty booklet for the Netherlands and the EU, and Voltwerk's own terms of "
                   "sale (edition 2025).", small),
         Paragraph("1. The OEM's pack warranty", h2),
         Paragraph("Each OEM warrants the traction-battery pack for a fixed number of years from the vehicle's date of "
                   "first admission (datum eerste toelating), wherever in the EU the vehicle was first admitted. The "
                   "warranty goes with the vehicle to every later holder and does not restart when the vehicle changes "
                   "country. None of the five carries a mileage limit on the pack.", body),
         _table(rows, [44, 38, 48, 34]),
         Spacer(1, 3 * mm),
         Paragraph("A pack the OEM replaces under its warranty is a new OEM pack and is covered for what remains of the "
                   "original period: the replacement does not start a new period. Once the period has ended the OEM "
                   "no longer replaces packs free of charge, and a failed pack is replaced at the owner's cost, in "
                   "practice by one of our refurbished packs.", body),
         Paragraph("2. Voltwerk's warranty on refurbished packs", h2),
         Paragraph("A Voltwerk refurbished pack carries our own warranty of %d years from the day it is fitted. A "
                   "refurbished pack that fails inside that period is replaced by another refurbished pack free of "
                   "charge; one that fails after it is replaced at the owner's cost. Either way the replacement comes "
                   "from the line." % P.VW_WARRANTY, body),
         Paragraph("3. Where the dates come from", h2),
         Paragraph("The date of first admission is the RDW's datum eerste toelating for the registration "
                   "(%s, field descriptions in %s). Fitting dates are in the replacement log (%s)."
                   % (F["register"], F["rdw_notes"], F["log"]), body)]
    _pdf(path, "Traction-battery warranty terms on the five platforms", "Warranty terms summary, February 2026", s,
         "Voltwerk Accu Revisie B.V. | Technical Services | warranty terms summary 2026")


def rdw_notes(path):
    body, h1, h2, small = _styles()
    fields = [
        ("kenteken", "Registration number, without dashes."),
        ("voertuigsoort", "Vehicle kind (Personenauto, Bedrijfsauto)."),
        ("merk", "Make as registered."),
        ("handelsbenaming", "Trade name as registered; varies by variant and over the years."),
        ("datum_eerste_toelating", "Date the vehicle was first admitted to the road, in the Netherlands or elsewhere "
                                   "(YYYYMMDD)."),
        ("datum_eerste_tenaamstelling_in_nederland", "Date the vehicle was first registered to a holder in the "
                                                     "Netherlands (YYYYMMDD)."),
        ("datum_tenaamstelling", "Date of the registration to the current holder (YYYYMMDD)."),
        ("vervaldatum_apk", "Date the current periodic inspection (APK) expires (YYYYMMDD)."),
        ("inrichting", "Body type."),
        ("eerste_kleur", "Main colour."),
        ("aantal_zitplaatsen", "Seats."),
        ("massa_rijklaar", "Mass in running order, kg."),
        ("catalogusprijs", "List price in euros where the RDW holds one; often empty for vehicles first admitted "
                           "abroad."),
        ("europese_voertuigcategorie", "EU vehicle category (M1, N1)."),
        ("taxi_indicator", "Registered as a taxi (Ja/Nee)."),
    ]
    rows = [["Field", "Description"]] + [[a, b] for a, b in fields]
    s = [Paragraph(P.ORG + " | Operations planning", small), Spacer(1, 2 * mm),
         Paragraph("RDW register extract of 30 June 2026: field descriptions", h1),
         Paragraph("Extract taken by operations planning from the RDW open data (Gekentekende voertuigen) on 1 July "
                   "2026, for the five platforms we refurbish.", small),
         Paragraph("What the extract holds", h2),
         Paragraph("One row per registration: every vehicle of the makes VOLTA, KESTREL, PLOVER, MIRA and TERN on the "
                   "Dutch register on 30 June 2026 that was insured and not reported exported on that date. Each make "
                   "is one of our platforms (VOLTA the Volta hatch, KESTREL the Kestrel SUV, PLOVER the Plover saloon, "
                   "MIRA the Mira city car, TERN the Tern van). Fields are as the RDW publishes them; dates are "
                   "YYYYMMDD.", body),
         _table(rows, [62, 112])]
    _pdf(path, "RDW register extract of 30 June 2026: field descriptions", "Field descriptions", s,
         "Voltwerk Accu Revisie B.V. | operations planning | RDW extract 2026-06-30")


def line_notes(path):
    F = P.FILES
    t = """# Line exports for 2025-26: what is in each file

Exports pulled on 1 July 2026 from the MES, the stores system and the end-of-line test benches, for the 2025-26
planning work. Times are local time (Europe/Amsterdam). Platform codes in the MES: VH Volta hatch, KS Kestrel SUV,
PS Plover saloon, MC Mira city car, TV Tern van.

## %(workorders)s

One row per work order opened from 1 December 2024: `werkorder`, `geopend` (date opened), `platformcode`, `kerncode`
(the core the pack is built on), `status` on 1 July 2026 (VERZONDEN shipped, AFGEKEURD scrapped, IN BEWERKING still on
the line) and `afgesloten` (date closed).

## %(stores)s

Stores movements against work orders with a booking date from 1 July 2025 to 30 June 2026. `soort` is UITGIFTE for an
issue and RETOUR for a return to stores, which carries a negative quantity. `aantal` is in the article's unit (see
%(parts)s). `partij` is the supplier lot the units came from. `picklijst` is the pick list the units went out on; one
list can serve several work orders.

## %(receipts)s

Goods received into stores from 1 July 2025 to 30 June 2026, by supplier lot, with the supplier's delivery note.

## %(parts)s

The module articles, current (actief) and superseded (vervallen): description, platform, stores unit, the modules in
one unit and the supplier.

## %(eol)s

Every end-of-line test run since the benches went live on 6 January 2025. `pakket_id` is the pack's serial number,
`programma` the test programme, `resultaat` PASS, FAIL or ABORT. ABORT means the bench stopped the run itself (a
bench, contact or cooling fault); an aborted run is no test of the pack, which is run again.

## %(bms)s

Every BMS paired to a pack on a work order opened from 1 December 2024: `soort` is MONTAGE when the pack is assembled
and VERVANGING when a BMS is replaced.

## %(shipments)s

Packs shipped to garages from 1 July 2025 to 30 June 2026, one row per shipment.

## %(intake)s

Cores received from 1 July 2025 to 30 June 2026 and graded at dock 1 or dock 2: `herkomst` is garage (the failed pack
sent back when a garage fits one of ours) or inkoop (cores bought in), `keuring` is the core's grade on 1 July 2026, A,
B or AFGEKEURD, with the reason.

## %(recycler)s

The recycler's waste transfer notes for its Wednesday collections from 2 July 2025 to 1 July 2026, one line per pack
collected, with the date on the quarantine tag.
""" % F
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(t)


def plan(path, created):
    F = P.FILES
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    p = d.add_paragraph()
    r = p.add_run("Voltwerk Accu Revisie B.V. | Operations planning")
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    d.add_heading("Tooling plan 2027: refurbishment slots by platform", level=1)
    p = d.add_paragraph()
    r = p.add_run("Concept 3, for the operations meeting of 2 September 2026. %s, operations planning, 14 August "
                  "2026. Fixtures and module stock are ordered by 30 September." % P.PEOPLE["planner"])
    r.italic = True
    r.font.size = Pt(9.5)
    d.add_heading("1. Capacity", level=2)
    d.add_paragraph("In 2027 the line has %s refurbishment slots on its four bays, two shifts a day. Fixtures, module "
                    "stock and the end-of-line test programmes are platform-specific and are committed before the year "
                    "starts, in lots of %d slots. A lot cannot be moved to another platform during the year: moving one "
                    "means another fixture set and a revalidated test programme, and the bay stands for about two "
                    "weeks while that is done." % ("{:,}".format(P.SLOTS), P.LOT))
    d.add_heading("2. How the slots are split", level=2)
    d.add_paragraph("The %s slots are split across the five platforms in proportion to the forecast of traction-battery "
                    "pack failures outside the OEM warranty on each platform in the Netherlands in calendar 2027. Lots "
                    "are set by the largest-remainder rule: each platform first takes the whole lots its share covers, "
                    "and the lots left over go to the platforms with the largest remainders." % "{:,}".format(P.SLOTS))
    d.add_paragraph("Garage partners order against the split through the year; a platform that runs out of lots is "
                    "served from the next year's tooling, so a short platform costs us its fittings, not just a delay.")
    d.add_heading("3. What counts as a failure", level=2)
    d.add_paragraph("A failure counts when a pack fails in 2027 on a car in use in the Netherlands and the OEM's pack "
                    "warranty on that car has ended (%s). A Voltwerk refurbished pack's failure counts whether or not "
                    "our own warranty still covers it: either way the line supplies the pack." % F["warranty"])
    d.add_paragraph("The forecast takes no car out of use: a car in use in the Netherlands in 2027 is taken to stay in "
                    "use for the rest of the year. Packs fail as the life curves say (%s), read by straight lines "
                    "between whole years of pack age. A pack's age runs from the day it was fitted, or from the car's "
                    "first admission for the pack the car was built with." % F["curves"])
    d.add_heading("4. Sources", level=2)
    for k, txt in [("register", "the RDW register extract of 30 June 2026 for the five platforms (field descriptions "
                                "in %s)" % F["rdw_notes"]),
                   ("log", "every traction-battery pack replacement on a car registered here up to 30 June 2026, from "
                           "the car's first admission: the OEMs' EU-wide dealer records, matched to the Dutch "
                           "registration by VIN, and our own fitting records"),
                   ("curves", "the share of packs failed by each whole year of pack age, from the OEMs' service "
                              "documentation, with our own curve for refurbished packs from our warranty returns "
                              "2019 to 2025"),
                   ("warranty", "the warranty terms")]:
        p = d.add_paragraph(style="List Bullet")
        r = p.add_run(F[k])
        r.bold = True
        p.add_run(": " + txt + ".")
    d.add_heading("5. What a lot carries", level=2)
    d.add_paragraph("Each lot of %d is ordered as one set: the module stock to refurbish %d packs, the bench time to "
                    "test them, re-runs included, and the cores to have on hand, rejects at intake included. The rates "
                    "come from the line's exports for 2025-26 (%s). Module lead times, minimum order quantities and "
                    "supplier prices are in %s; the Kestrel SUV's carriers come from one supplier, so its lots need the "
                    "order in first." % (P.LOT, P.LOT, F["line_notes"], F["leadtimes"]))
    d.add_heading("6. Still to settle", level=2)
    d.add_paragraph("For the meeting of 2 September: the slots and the forecast per platform, and what each lot "
                    "carries, for sign-off by the head of operations. %s to price the module orders and to put the "
                    "year's revenue against the garage prices (%s)." % (P.PEOPLE["controller"], F["prices"]))
    d.core_properties.author = P.PEOPLE["planner"]
    d.core_properties.title = "Tooling plan 2027"
    d.save(path)


def prices(path):
    body, h1, h2, small = _styles()
    rs = np.random.default_rng([P.SEED_DOCS, 1])
    base = {"Volta hatch": 4865, "Kestrel SUV": 7415, "Plover saloon": 6245, "Mira city car": 3195, "Tern van": 7960}
    rows = [["Platform", "Refurbished pack, excl. VAT", "Core deposit", "Fitting (garage partners)", "Lead time"]]
    for n in P.NAMES:
        rows.append([n, "EUR %s" % "{:,}".format(base[n]), "EUR %s" % "{:,}".format(int(round(base[n] * 0.18 / 5.0) * 5)),
                     "EUR %d" % int(305 + 20 * rs.integers(0, 5)), "%d working days" % (3 + int(rs.integers(0, 4)))])
    other = [["Other charges", "Excl. VAT"],
             ["Transport to the garage, per pack", "EUR 48.50"],
             ["Next-working-day delivery, per pack", "EUR 95.00"],
             ["Core collected without a replacement order", "EUR 62.00"],
             ["Diagnostic report on a returned pack", "EUR 37.50"]]
    s = [Paragraph(P.ORG + " | Sales", small), Spacer(1, 2 * mm),
         Paragraph("Refurbished traction-battery packs: price list 2026", h1),
         Paragraph("Valid from 1 January 2026 for garage partners in the Netherlands. Prices exclude VAT.", small),
         _table(rows, [36, 40, 28, 38, 30]),
         Spacer(1, 3 * mm),
         _table(other, [96, 40]),
         Spacer(1, 3 * mm),
         Paragraph("The core deposit is refunded when the failed pack comes back to us complete and undamaged; a core "
                   "we reject at intake is not refunded. Packs carry our two-year warranty from fitting. Prices are "
                   "reviewed each January.", body)]
    _pdf(path, "Refurbished traction-battery packs: price list 2026", "Price list 2026", s,
         "Voltwerk Accu Revisie B.V. | price list 2026")


def leadtimes(path, ws_writer):
    rs = np.random.default_rng([P.SEED_DOCS, 2])
    import line as Ln
    rows = []
    for n in P.NAMES:
        rows.append([Ln.PART[n], Ln.SUPPLIER[n], n, int(6 + rs.integers(0, 9)), int([50, 100, 200][rs.integers(0, 3)]),
                     round(float(rs.uniform(180, 420)), 2)])
    for n in P.NAMES:
        rows.append(["BMS-%s-%d" % (Ln.BMS_PRE[n], 3 + int(rs.integers(0, 3))), Ln.SUPPLIER[n], n,
                     int(4 + rs.integers(0, 7)), int([20, 25, 50][rs.integers(0, 3)]),
                     round(float(rs.uniform(140, 260)), 2)])
    rows.append(["AFD-SET-UNI", "Hanseatic Module Supply", "alle", 3, 100, round(float(rs.uniform(11, 19)), 2)])
    rows.append(["KLP-KST-2", "Hanseatic Module Supply", "Kestrel SUV", int(8 + rs.integers(0, 4)), 40,
                 round(float(rs.uniform(30, 55)), 2)])
    ws_writer(path, rows)
