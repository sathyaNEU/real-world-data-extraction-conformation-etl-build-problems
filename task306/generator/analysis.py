"""The golden and every rung, computed from the shipped files only (target/), never from the generator's truth.

The golden: link each storm of record to the model's hindcast of it by date, read the record's storms in their
season's order, cost each simulated storm in the year loss table at what the record shows for a storm of its class
(a storm that strikes within weeks of the one before it, or a storm on a quiet season), and read the capital policy's
1-in-200 gross, net of the occurrence cover, and net of both covers."""
import csv
import datetime as dt
import os
from collections import Counter, defaultdict

import numpy as np
from openpyxl import load_workbook

import capital as C
import world as W

FOLLOW_WINDOW = 42      # the golden's reading of "within weeks": any window from 25 to 69 days classes the same storms


def _rows(target, name):
    with open(os.path.join(target, name), newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def _files():
    import pack as P
    return P.FILES


def load_ylt(target):
    F = _files()
    rows = _rows(target, F["ylt"])
    years = np.array([int(r["year"]) for r in rows])
    loss = np.array([float(r["loss_gbp"]) for r in rows])
    peak = [dt.datetime.strptime(r["peak_utc"], "%Y-%m-%dT%H:%MZ") for r in rows]
    return years, loss, peak


def ylt_follows(years, peak, window=FOLLOW_WINDOW):
    """a simulated storm follows another when it strikes within `window` days of the storm before it in its year"""
    out = np.zeros(len(years), dtype=bool)
    for i in range(1, len(years)):
        if years[i] == years[i - 1] and (peak[i].date() - peak[i - 1].date()).days <= window:
            out[i] = True
    return out


def record(target):
    """the record's storms linked to the hindcast: (date, hindcast loss, incurred, storm name)"""
    F = _files()
    wb = load_workbook(os.path.join(target, F["experience"]), read_only=True, data_only=True)
    ws = wb["Storms"]
    exp = []
    for r in list(ws.iter_rows(values_only=True))[1:]:
        d = r[2].date() if isinstance(r[2], dt.datetime) else r[2]
        exp.append({"ref": r[0], "name": r[1], "date": d, "claims": r[3], "incurred": float(r[4])})
    hind = [{"id": h["event_id"], "date": dt.datetime.strptime(h["peak_utc"], "%Y-%m-%dT%H:%MZ").date(),
             "loss": float(h["loss_gbp"])} for h in _rows(target, F["hindcast"])]
    linked = []
    for e in exp:
        m = [h for h in hind if 0 <= (h["date"] - e["date"]).days <= 1]
        assert len(m) == 1, (e, m)
        linked.append((e["date"], m[0]["loss"], e["incurred"], e["name"]))
    linked.sort()
    return linked


def record_follows(linked, window=FOLLOW_WINDOW):
    out = []
    for i, r in enumerate(linked):
        out.append(i > 0 and (r[0] - linked[i - 1][0]).days <= window)
    return np.array(out)


def loads(target, how="golden", window=FOLLOW_WINDOW, est="rt"):
    """(load per simulated storm, description) under each reading of the record"""
    years, loss, peak = load_ylt(target)
    rec = record(target)
    hind = np.array([r[1] for r in rec])
    act = np.array([r[2] for r in rec])

    def ratio(m):
        if est == "rt":
            return act[m].sum() / hind[m].sum()
        if est == "mr":
            return float((act[m] / hind[m]).mean())
        return float(np.median(act[m] / hind[m]))
    if how == "raw":
        return np.ones(len(loss)), 1.0
    if how == "pooled":
        p = ratio(np.ones(len(rec), dtype=bool))
        return np.full(len(loss), p), p
    if how == "size":
        b = np.polyfit(np.log(hind), act / hind, 1)
        return np.polyval(b, np.log(loss)), tuple(b)
    if how == "year":
        rc = Counter(r[0].year for r in rec)
        rm = np.array([rc[r[0].year] > 1 for r in rec])
        tc = Counter(years.tolist())
        tm = np.array([tc[y] > 1 for y in years])
        return np.where(tm, ratio(rm), ratio(~rm)), (ratio(~rm), ratio(rm))
    if how == "rank":
        seen = defaultdict(int)
        rr = []
        for r in rec:
            rr.append(seen[r[0].year])
            seen[r[0].year] += 1
        rr = np.array(rr)
        tr = np.zeros(len(years), dtype=int)
        tr[1:] = np.where(years[1:] == years[:-1], 1, 0)
        later = np.array([i > 0 and years[i] == years[i - 1] for i in range(len(years))])
        return np.where(later, ratio(rr > 0), ratio(rr == 0)), (ratio(rr == 0), ratio(rr > 0))
    rf = record_follows(rec, window)
    tf = ylt_follows(years, peak, window)
    if how == "allf":
        return np.full(len(loss), ratio(rf)), ratio(rf)
    return np.where(tf, ratio(rf), ratio(~rf)), (ratio(~rf), ratio(rf))


def programme(target, how="golden", **kw):
    years, loss, peak = load_ylt(target)
    ld, desc = loads(target, how, **kw)
    out = C.programme(years, loss, ld, W.BASE, W.SEC, W.N_YEARS, W.RANK_1IN200)
    out["value"] = C.value(out, W.COC, W.SEC)
    out["loads"] = desc
    return out


RUNGS = ["raw", "pooled", "size", "year", "rank", "allf", "golden"]


# ---------------------------------------------------------------------------------------------------------- ask A
def _d(s):
    return dt.datetime.strptime(s, "%d/%m/%Y").date()


def storm_lines(target, dating="pair"):
    """the 2026 windstorm claim lines (day of damage, incurred, line of business, claim number). The golden dates
    each line by the loss adjusters' day of damage, a claim's lines taken in the order they were notified against its
    instructions in the order they were given (a reopened claim with a second instruction is further damage done on
    the second instruction's day); the other datings are the rungs: every line at the claim's first instruction
    ("first") or its last ("last"), the instruction joined on the full claim number so a reopened line keeps its
    recorded date ("full"), the recorded loss date ("recorded") or the notified date ("notified")."""
    F = _files()
    lines = [c for c in _rows(target, F["claims"]) if c["peril"] == "Storm"]
    inst = defaultdict(list)
    for a in _rows(target, F["adjusters"]):
        inst[a["insurer_ref"].replace(" ", "")].append((a["instructed"], a["job_ref"], dt.date.fromisoformat(a["date_of_damage"])))
    for k in inst:
        inst[k].sort()
    by_claim = defaultdict(list)
    for c in lines:
        by_claim[c["claim_number"].split("/")[0]].append(c)
    out = []
    for base, cl in by_claim.items():
        cl.sort(key=lambda c: (_d(c["notified_date"]), c["claim_number"]))
        ins = inst.get(base, [])
        for k, c in enumerate(cl):
            rec = _d(c["loss_date"])
            if dating == "recorded" or not ins and dating != "notified":
                day = rec
            elif dating == "notified":
                day = _d(c["notified_date"])
            elif dating == "first":
                day = ins[0][2]
            elif dating == "last":
                day = ins[-1][2]
            elif dating == "full":
                day = ins[0][2] if "/" not in c["claim_number"] else rec
            else:
                day = ins[min(k, len(ins) - 1)][2]
            out.append((day, float(c["paid_gbp"]) + float(c["outstanding_gbp"]), c["line"], c["claim_number"]))
    return out


def occurrences(lines, floor=1_000_000.0, choose="best"):
    """the 72-hour loss occurrences over the windstorm days. "best": the Reassured's choice the wording gives it,
    over each run of consecutive days the split into periods of up to three days that maximises the expiring cover's
    recoveries (fewest periods on a tie); "greedy": each period starting on the first day not yet in one. Each
    occurrence reaching the floor is listed as (first day, last day, incurred, home, farm)."""
    by_day = defaultdict(float)
    by_line = defaultdict(lambda: defaultdict(float))
    for d, v, ln, _ in lines:
        by_day[d] += v
        by_line[d][ln] += v
    days = sorted(by_day)
    runs, cur = [], []
    for d in days:
        if cur and (d - cur[-1]).days > 1:
            runs.append(cur)
            cur = []
        cur.append(d)
    if cur:
        runs.append(cur)
    periods = []
    for run in runs:
        if choose == "greedy":
            i = 0
            while i < len(run):
                j = i
                while j + 1 < len(run) and (run[j + 1] - run[i]).days <= 2:
                    j += 1
                periods.append(run[i:j + 1])
                i = j + 1
            continue
        n = len(run)
        best = [None] * (n + 1)
        best[0] = (0.0, [])
        for i in range(1, n + 1):
            cand = []
            for k in (1, 2, 3):
                if i - k >= 0 and best[i - k] is not None and (run[i - 1] - run[i - k]).days <= 2:
                    tot = sum(by_day[x] for x in run[i - k:i])
                    r = min(max(tot - W.EXP["retention"], 0.0), W.EXP["limit"])
                    cand.append((best[i - k][0] + r, best[i - k][1] + [run[i - k:i]]))
            best[i] = max(cand, key=lambda x: (round(x[0], 2), -len(x[1])))
        periods.extend(best[n][1])
    out = []
    for p in periods:
        tot = sum(by_day[x] for x in p)
        if tot >= floor:
            out.append((p[0], p[-1], tot, sum(by_line[x]["Home"] for x in p), sum(by_line[x]["Farm"] for x in p)))
    return sorted(out)


def ask_a(target, dating="pair", choose="best"):
    """each 2026 loss occurrence of a million pounds or more, and what the expiring cover recovers on the year and the
    reinstatement premium it calls (one reinstatement, 100 per cent of the annual premium pro rata as to amount)"""
    occ = occurrences(storm_lines(target, dating), choose=choose)
    e = W.EXP
    rows, used, rp_tot = [], 0.0, 0.0
    for start, end, tot, home, farm in occ:
        rec = min(max(tot - e["retention"], 0.0), e["limit"], e["limit"] * (1 + e["reinstatements"]) - used)
        reinst_room = max(e["limit"] * e["reinstatements"] - used, 0.0)
        rp = e["premium"] * min(rec, reinst_room) / e["limit"]
        used += rec
        rp_tot += rp
        rows.append({"start": start, "end": end, "incurred": tot, "home": home, "farm": farm, "recovery": rec,
                     "reinstatement": rp})
    return {"occurrences": rows, "recovery": used, "reinstatement": rp_tot}


# ---------------------------------------------------------------------------------------------------------- ask B
def ask_b(target, shares="golden", county="holding"):
    """the farm book's own sums insured at the extract date by the county each holding stands in, buildings and farm
    contents apart, in pounds. The golden takes each coinsured policy's share from the register's row for its main
    holding, under the holding's current CPH or the one it had before it was renumbered, and applies it to every
    holding on the policy; the rungs: the register matched to the current main holding only ("main"), matched holding
    by holding ("holding"), or not applied ("none")."""
    F = _files()
    wb = load_workbook(os.path.join(target, F["farm"]), read_only=True, data_only=True)
    pols = list(wb["Policies"].iter_rows(values_only=True))[1:]
    locs = list(wb["Holdings"].iter_rows(values_only=True))[1:]
    reg = {c["holding_cph"]: float(c["hartcombe_share"].rstrip("%")) / 100.0 for c in _rows(target, F["coins"])}
    main_prev = {}
    for l in locs:
        if l[1] == 1 and l[3]:
            main_prev[l[0]] = l[3].replace("/", "")
    pol_share, pol_county = {}, {}
    for p in pols:
        key = p[1].replace("/", "")
        sh = reg.get(key)
        if sh is None and shares == "golden":
            sh = reg.get(main_prev.get(p[0], ""))
        pol_share[p[0]] = 1.0 if (sh is None or shares in ("none", "holding")) else sh
        pol_county[p[0]] = p[2]
    out = {c: [0.0, 0.0] for c in W.COUNTIES}
    for l in locs:
        sh = pol_share[l[0]]
        if shares == "holding":
            sh = reg.get(l[2].replace("/", ""), 1.0)
        cty = l[4] if county == "holding" else pol_county[l[0]]
        out[cty][0] += float(l[5]) * 1000.0 * sh
        out[cty][1] += float(l[6]) * 1000.0 * sh
    return {c: tuple(v) for c, v in out.items()}
