"""Cut the task306 pack from the world and the line."""
import csv
import datetime as dt
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import docs
import line as Ln
import params as P
import postprocess
import records as R
import world as Wd

FILES = P.FILES
PACK = list(FILES)
DISTRACTORS = ("prices", "leadtimes")
CONTEXT = ("plan",)
LABEL = {"Volta hatch": ["VOLTA HATCH", "VOLTA E-HATCH"], "Kestrel SUV": ["KESTREL SUV", "KESTREL EV"],
         "Plover saloon": ["PLOVER SALOON", "PLOVER E"], "Mira city car": ["MIRA CITY", "MIRA E"],
         "Tern van": ["TERN VAN", "TERN E-VAN"]}


def _csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)


def _xlsx(path, sheets, made, saved, author):
    wb = Workbook()
    for i, (title, header, rows, widths, fmts) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDE5EA")
            c.alignment = Alignment(vertical="top", wrap_text=True)
        for r in rows:
            ws.append(list(r))
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        for col, f in (fmts or {}).items():
            j = header.index(col) + 1
            for k in range(2, len(rows) + 2):
                ws.cell(row=k, column=j).number_format = f
        ws.freeze_panes = "A2"
    wb.properties.creator = author
    wb.properties.lastModifiedBy = author
    wb.save(path)
    postprocess.scrub(path, P.ORG, saved[:10])
    postprocess.finish_office(path, made, saved, P.ORG)


def ts(t):
    return t.strftime("%Y-%m-%d %H:%M")


def build(target):
    os.makedirs(target, exist_ok=True)
    F = {k: os.path.join(target, v) for k, v in FILES.items()}
    W, ev = Wd.world()
    plate = R.plates(W)
    reg, _ = R.register(W, plate)
    reg.to_csv(F["register"], index=False, lineterminator="\n")
    log = R.replacement_log(W, ev, plate)
    log.to_csv(F["log"], index=False, lineterminator="\n")
    cv = R.curves()
    _csv(F["curves"], list(cv.columns),
         [[int(r[0])] + ["%.4f" % x for x in r[1:]] for r in cv.itertuples(index=False)])
    L = Ln.Line(W, ev).build()
    fy = lambda d: P.FY0 <= d <= P.FY1
    rs = __import__("numpy").random.default_rng([P.SEED_LINE, 9])
    # MES work orders
    rows = []
    for w in sorted(L.wo, key=lambda w: (w["geopend"], w["werkorder"])):
        if w["geopend"] < Ln.WO_FROM or w["geopend"] >= Ln.EXPORT:
            continue
        closed = w["afgesloten"] if (w["afgesloten"] and w["afgesloten"] < Ln.EXPORT) else None
        rows.append([w["werkorder"], w["geopend"].isoformat(), P.CODE[w["platform"]], w["kerncode"],
                     w["status_export"], closed.isoformat() if closed else ""])
    _csv(F["workorders"], ["werkorder", "geopend", "platformcode", "kerncode", "status", "afgesloten"], rows)
    # stores movements
    qs = sorted((q for q in L.reqs if fy(q["t"].date())), key=lambda q: (q["t"], q["pick"] or "", q["werkorder"]))
    rows = [[str(1840000 + i), ts(q["t"]), q["soort"], q["artikel"], q["partij"], q["aantal"], q["werkorder"],
             q["pick"]] for i, q in enumerate(qs, 1)]
    _csv(F["stores"], ["mutatie_id", "geboekt_op", "soort", "artikelnummer", "partij", "aantal", "werkorder",
                       "picklijst"], rows)
    # goods received
    rec = sorted((x for x in L.receipts if fy(x["datum"])), key=lambda x: (x["datum"], x["partij"]))
    _csv(F["receipts"], ["ontvangst", "datum", "artikelnummer", "partij", "aantal", "leverancier", "pakbon"],
         [["ONT%06d" % (71200 + 3 * i), x["datum"].isoformat(), x["artikel"], x["partij"], x["aantal"],
           x["leverancier"], x["pakbon"]] for i, x in enumerate(rec, 1)])
    # parts
    desc = {"Volta hatch": "Module 12S Volta hatch (A-grade, gereviseerd)",
            "Kestrel SUV": "Moduledrager Kestrel SUV, 2 modules",
            "Plover saloon": "Module 14S Plover saloon",
            "Mira city car": "Module 8C Mira city car",
            "Tern van": "Module 18P Tern van"}
    unit = {"Kestrel SUV": "drager"}
    arts = [[Ln.PART[n], desc[n], n, unit.get(n, "stuk"), P.PER_UNIT[n], Ln.SUPPLIER[n], "actief"] for n in P.NAMES]
    arts += [["MOD-VLT-12", "Module 12S Volta hatch", "Volta hatch", "stuk", 1, "Lucanor Cells GmbH", "vervallen"],
             ["MOD-KST-1", "Module Kestrel SUV, los (reparatie drager)", "Kestrel SUV", "stuk", 1,
              "Hanseatic Module Supply", "actief"],
             ["MOD-PLV-14", "Module 14S Plover saloon (eerste serie)", "Plover saloon", "stuk", 1, "Lucanor Cells GmbH",
              "vervallen"],
             ["MOD-MRA-8B", "Module 8B Mira city car", "Mira city car", "stuk", 1, "Celeste Batteries S.r.l.",
              "vervallen"]]
    arts.sort(key=lambda a: (P.NAMES.index(a[2]), a[6] != "actief", a[0]))
    _xlsx(F["parts"], [("Artikelen", ["artikelnummer", "omschrijving", "platform", "eenheid", "modules_per_eenheid",
                                      "leverancier", "status"],
                        arts, [16, 42, 16, 10, 20, 26, 11], None)],
          "2025-03-03T10:12:00Z", "2026-02-11T08:47:00Z", "Stores")
    # end-of-line runs
    runs = [x for x in L.runs if x["start"].date() >= Ln.GO_LIVE and x["start"].date() <= P.FY1]
    _csv(F["eol"], ["run_id", "bank", "pakket_id", "programma", "start", "einde", "resultaat"],
         [[x["run"], x["bank"], x["pakket"], x["programma"], ts(x["start"]), ts(x["einde"]), x["resultaat"]]
          for x in runs])
    # BMS pairings
    opened = {w["werkorder"]: w["geopend"] for w in L.wo}
    bms = sorted((b for b in L.bms if opened[b["werkorder"]] >= Ln.WO_FROM and b["t"].date() < Ln.EXPORT),
                 key=lambda b: (b["t"], b["werkorder"]))
    _csv(F["bms"], ["werkorder", "bms_serienummer", "gekoppeld_op", "soort"],
         [[b["werkorder"], b["bms"], ts(b["t"]), b["soort"]] for b in bms])
    # shipments
    sh = sorted((s for s in L.ships if fy(s["verzonden"])), key=lambda s: (s["verzonden"], s["zending"]))
    _csv(F["shipments"], ["zending", "verzonden_op", "werkorder", "pakket_id", "garage"],
         [[s["zending"], s["verzonden"].isoformat(), s["werkorder"], s["pakket"], "G%04d" % s["garage"]] for s in sh])
    # core intake, each core with its grade on the morning of the export
    ins = sorted((x for x in L.intake if x["scanned"] and fy(x["ontvangen"].date())), key=lambda x: x["ontvangen"])
    _csv(F["intake"], ["kerncode", "ontvangen_op", "dock", "label_platform", "herkomst", "keuring", "reden_afkeur"],
         [[x["kerncode"], ts(x["ontvangen"]), x["dock"], LABEL[x["platform"]][int(rs.integers(0, 2))], x["herkomst"],
           x["keuring_export"], x["reden_export"]] for x in ins])
    # the recycler's notes
    rr = [x for x in L.recycled if dt.date(2025, 7, 2) <= x["ophaal"] <= Ln.EXPORT]
    _xlsx(F["recycler"], [("Begeleidingsbrieven", ["begeleidingsbrief", "ophaaldatum", "regel", "un_nummer",
                                                   "omschrijving", "bruto_kg", "kerncode", "datum_quarantaine"],
                           [[x["note"], x["ophaal"], x["regel"], "UN3480",
                             "Lithium-ion tractiebatterij, %s" % x["platform"], x["kg"], x["kerncode"],
                             x["quarantaine"]] for x in rr],
                           [18, 13, 7, 11, 40, 10, 12, 18], {"ophaaldatum": "dd-mm-yyyy",
                                                             "datum_quarantaine": "dd-mm-yyyy"})],
          "2025-07-03T07:40:00Z", "2026-07-01T13:05:00Z", P.ORG)
    docs.line_notes(F["line_notes"])
    # documents
    docs.warranty(F["warranty"])
    postprocess.scrub(F["warranty"], "Microsoft Word 2019", "2026-02-03T11:20:00")
    docs.rdw_notes(F["rdw_notes"])
    postprocess.scrub(F["rdw_notes"], "Microsoft Word 2019", "2026-07-01T14:02:00")
    docs.plan(F["plan"], None)
    postprocess.scrub(F["plan"], P.ORG, "2026-08-14")
    postprocess.finish_office(F["plan"], "2026-08-06T07:58:00Z", "2026-08-14T15:31:00Z", P.ORG,
                              person=P.PEOPLE["planner"], revision=9)
    docs.prices(F["prices"])
    postprocess.scrub(F["prices"], "Microsoft Word 2019", "2025-12-18T09:00:00")

    def lt(path, rows):
        _xlsx(path, [("Levertijden", ["artikelnummer", "leverancier", "platform", "levertijd_weken", "minimale_afname",
                                      "prijs_per_eenheid_eur"], rows, [16, 26, 16, 16, 16, 20],
                      {"prijs_per_eenheid_eur": "#,##0.00"})],
              "2026-01-12T09:30:00Z", "2026-05-20T14:12:00Z", "Purchasing")
    docs.leadtimes(F["leadtimes"], lt)
    return W, ev, L
