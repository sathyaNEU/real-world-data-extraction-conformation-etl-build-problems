"""Independent verifier for task306: reads only target.zip, on its own code path (no generator module), takes the
covers' terms from the quotation and the capital policy's rules from its text, recomputes the call and every graded
figure, and compares them with generator/figures.json.

    python3 task306/generator/verify.py [target.zip]
"""
import csv
import datetime as dt
import email
import io
import json
import os
import re
import sys
import zipfile
from collections import defaultdict

import numpy as np
from openpyxl import load_workbook
from pdfminer.high_level import extract_text

HERE = os.path.dirname(os.path.abspath(__file__))
BAD = []


def check(label, cond, detail=""):
    print("%s  %s  %s" % ("PASS" if cond else "FAIL", label, detail))
    if not cond:
        BAD.append(label)


def text(z, name):
    return re.sub(r"\s+", " ", extract_text(io.BytesIO(z.read(name))))


def rows(z, name):
    return list(csv.DictReader(io.StringIO(z.read(name).decode("utf-8"))))


def gbp(s):
    return float(s.replace(",", ""))


def nets(year, loss, n, base, sec):
    """each simulated year's loss gross, net of the occurrence cover, and net of both, in date order within year"""
    gross = np.zeros(n)
    rec_b = np.zeros(n)
    rec_s = np.zeros(n)
    seen = set()
    for y, l in zip(year, loss):
        gross[y - 1] += l
        rec_b[y - 1] += min(max(l - base["ret"], 0.0), base["lim"])
        if y in seen:
            rec_s[y - 1] += min(max(min(l, base["ret"]) - sec["ret"], 0.0), sec["lim"])
        seen.add(y)
    rec_b = np.minimum(rec_b, 2 * base["lim"])
    rec_s = np.minimum(rec_s, 2 * sec["lim"])
    rp_b = base["prem"] * np.minimum(rec_b, base["lim"]) / base["lim"]
    rp_s = sec["prem"] * np.minimum(rec_s, sec["lim"]) / sec["lim"]
    nb = gross - rec_b + rp_b
    return gross, nb, nb - rec_s + rp_s, rec_s, rp_s


def tvar(x, k):
    return float(np.sort(x)[-k:].mean())


def main(zp):
    z = zipfile.ZipFile(zp)
    names = set(z.namelist())
    fig = json.load(open(os.path.join(HERE, "figures.json")))
    cp = text(z, "CP-04_capital_policy_windstorm_2026.pdf")
    qt = text(z, "penwarden_windstorm_quotation_2027.pdf")
    wd = text(z, "windstorm_xl_wording_2026.pdf")
    check("pdfminer reads the three documents (positive control)", "Capital Policy CP-04" in cp
          and "Second-event cover" in qt and "72 consecutive hours" in wd)
    k = int(re.search(r"average of the (\d+) worst of the ([\d,]+) simulated years", cp).group(1))
    nyears = int(gbp(re.search(r"average of the (\d+) worst of the ([\d,]+) simulated years", cp).group(2)))
    coc = float(re.search(r"cost of capital of (\d+) per cent", cp).group(1)) / 100
    terms = re.findall(r"GBP ([\d,]+) each and every Loss Occurrence excess of GBP ([\d,]+)", qt)
    prems = re.findall(r"GBP ([\d,]+), payable in", qt)
    base = {"lim": gbp(terms[0][0]), "ret": gbp(terms[0][1]), "prem": gbp(prems[0])}
    sec = {"lim": gbp(terms[1][0]), "ret": gbp(terms[1][1]), "prem": gbp(prems[1])}
    check("the covers' terms read from the quotation", base["ret"] > sec["ret"] and sec["lim"] < base["lim"],
          "A %s, B %s" % (base, sec))
    # the record linked to the hindcast by date
    wb = load_workbook(io.BytesIO(z.read("windstorm_experience_2001-2025.xlsx")), read_only=True, data_only=True)
    exp = [(r[2].date() if isinstance(r[2], dt.datetime) else r[2], float(r[4]))
           for r in list(wb["Storms"].iter_rows(values_only=True))[1:]]
    hc = [(dt.datetime.strptime(r["peak_utc"], "%Y-%m-%dT%H:%MZ").date(), float(r["loss_gbp"]))
          for r in rows(z, "seaborne_windstorm_hindcast_2001-2025.csv")]
    link = []
    for d, inc in sorted(exp):
        m = [h for h in hc if 0 <= (h[0] - d).days <= 1]
        link.append((d, m[0][1] if len(m) == 1 else None, inc))
    check("every storm of record links to one hindcast storm", all(l[1] for l in link), "%d storms" % len(link))
    fol = [i > 0 and (link[i][0] - link[i - 1][0]).days <= 42 for i in range(len(link))]
    h = np.array([l[1] for l in link]); a = np.array([l[2] for l in link]); f = np.array(fol)
    lf, ll = a[f].sum() / h[f].sum(), a[~f].sum() / h[~f].sum()
    check("the two ratios match the write-up", round(lf, 2) == fig["record"]["load_follow"]
          and round(ll, 2) == fig["record"]["load_other"], "%.3f and %.3f" % (lf, ll))
    # the table, each storm classed against the one before it in its simulated year
    yl = rows(z, "seaborne_windstorm_ylt_2027_portfolio.csv")
    year = np.array([int(r["year"]) for r in yl]); loss = np.array([float(r["loss_gbp"]) for r in yl])
    peak = [dt.datetime.strptime(r["peak_utc"], "%Y-%m-%dT%H:%MZ").date() for r in yl]
    tcls = np.array([i > 0 and year[i] == year[i - 1] and (peak[i] - peak[i - 1]).days <= 42 for i in range(len(yl))])
    costed = loss * np.where(tcls, lf, ll)
    gross, nb, nboth, rec_s, rp_s = nets(year, costed, nyears, base, sec)
    G, NB, NBO = tvar(gross, k), tvar(nb, k), tvar(nboth, k)
    rel = NB - NBO
    margin = sec["prem"] + rp_s.mean() - rec_s.mean()
    value = coc * rel - margin
    cap = fig["capital_gbp_m"]
    check("the 1-in-200 losses match", (round(G / 1e6), round(NB / 1e6), round(NBO / 1e6)) ==
          (cap["gross"], cap["net_occurrence"], cap["net_both"]), "%.2f %.2f %.2f" % (G / 1e6, NB / 1e6, NBO / 1e6))
    check("the release matches", round(rel / 1e6) == cap["release"], "%.2fm" % (rel / 1e6))
    check("the margin and the worth match", round(margin / 1e6, 1) == fig["cover_gbp_m"]["margin"]
          and round(value / 1e6, 1) == fig["cover_gbp_m"]["value"], "%.3f and %+.3f" % (margin / 1e6, value / 1e6))
    check("the call: the cover adds value", value > 0 and fig["call"] == "add the second-event cover")
    pooled = a.sum() / h.sum()
    _, pnb, pnbo, prs, prp = nets(year, loss * pooled, nyears, base, sec)
    pv = coc * (tvar(pnb, k) - tvar(pnbo, k)) - (sec["prem"] + prp.mean() - prs.mean())
    check("the default path (one ratio for every storm) values the cover below nil", pv < 0, "%+.3fm" % (pv / 1e6))
    # ask A on the claims and the adjusters' instructions
    inst = defaultdict(list)
    for r in rows(z, "loss_adjuster_instructions_2026.csv"):
        inst[r["insurer_ref"].replace(" ", "")].append((r["instructed"], r["job_ref"], r["date_of_damage"]))
    lines = defaultdict(list)
    for c in rows(z, "claims_extract_2026-10-31.csv"):
        if c["peril"] == "Storm":
            lines[c["claim_number"].split("/")[0]].append(c)
    day_amt = defaultdict(lambda: [0.0, 0.0])
    for ref, cl in lines.items():
        cl.sort(key=lambda c: (dt.datetime.strptime(c["notified_date"], "%d/%m/%Y"), c["claim_number"]))
        ins = sorted(inst.get(ref, []))
        for i, c in enumerate(cl):
            d = dt.date.fromisoformat(ins[min(i, len(ins) - 1)][2]) if ins else dt.datetime.strptime(c["loss_date"], "%d/%m/%Y").date()
            day_amt[d][0 if c["line"] == "Home" else 1] += float(c["paid_gbp"]) + float(c["outstanding_gbp"])
    days = sorted(day_amt)
    occ, i = [], 0
    while i < len(days):
        j = i
        while j + 1 < len(days) and (days[j + 1] - days[i]).days <= 2 and (days[j + 1] - days[j]).days <= 1:
            j += 1
        tot = [sum(day_amt[d][q] for d in days[i:j + 1]) for q in (0, 1)]
        if sum(tot) >= 1e6:
            occ.append((days[i], tot))
        i = j + 1
    ret, lim, prem = gbp(re.search(r"excess of GBP ([\d,]+) each and every", wd).group(1)), \
        gbp(re.search(r"up to GBP ([\d,]+) each and every", wd).group(1)), \
        gbp(re.search(r"annual premium of GBP ([\d,]+)", wd).group(1))
    recov = sum(min(max(sum(t) - ret, 0), lim) for _, t in occ)
    rp = prem * min(recov, lim) / lim
    got = [(round(t[0], -3), round(t[1], -3)) for _, t in occ]
    want = [(o["home"], o["farm"]) for o in fig["occurrences_2026"]]
    check("ask A: each occurrence's home and farm cost", got == [tuple(map(float, w)) for w in want], str(got))
    check("ask A: the recovery and the reinstatement premium", (round(recov, -3), round(rp, -3)) ==
          (fig["recovery_2026"], fig["reinstatement_2026"]), "%.0f and %.0f" % (recov, rp))
    # ask B on the farm book and the coinsurance register
    fw = load_workbook(io.BytesIO(z.read("farm_policies_and_locations_2026-10-31.xlsx")), read_only=True, data_only=True)
    pols = list(fw["Policies"].iter_rows(values_only=True))[1:]
    hold = list(fw["Holdings"].iter_rows(values_only=True))[1:]
    reg = {r["holding_cph"]: float(r["hartcombe_share"].rstrip("%")) / 100 for r in rows(z, "coinsurance_register_2026.csv")}
    keys = defaultdict(set)
    for p in pols:
        keys[p[0]].add(p[1].replace("/", ""))
    for l in hold:
        if l[1] == 1 and l[3]:
            keys[l[0]].add(l[3].replace("/", ""))
    share = {p: next((reg[x] for x in ks if x in reg), 1.0) for p, ks in keys.items()}
    si = defaultdict(lambda: [0.0, 0.0])
    for l in hold:
        si[l[4]][0] += l[5] * 1000 * share[l[0]]
        si[l[4]][1] += l[6] * 1000 * share[l[0]]
    ok_b = all(round(si[c][0] / 1e6, 1) == v["buildings"] and round(si[c][1] / 1e6, 1) == v["farm_contents"]
               for c, v in fig["farm_si_gbp_m"].items())
    check("ask B: each county's buildings and farm contents, our share", ok_b,
          "; ".join("%s %.1f/%.1f" % (c, v[0] / 1e6, v[1] / 1e6) for c, v in sorted(si.items())))
    m = email.message_from_bytes(z.read("2027_windstorm_programme.eml"))
    check("the request email reads back", "second-event" in m.get_payload())
    check("every pack file is present and nothing else", len(names) == 14, str(len(names)))
    print("verify: %d failed" % len(BAD) if BAD else "verify: all passed")
    return len(BAD)


if __name__ == "__main__":
    zp = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(HERE), "target.zip")
    sys.exit(1 if main(zp) else 0)
