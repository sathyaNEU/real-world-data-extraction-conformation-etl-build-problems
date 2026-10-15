"""Independent verifier for task261: reads only the shipped bundle and the golden workbook, recomputes every graded
figure on its own code path (its own weighted median, its own replicate and implicate arithmetic, the asks by table
joins), and compares with the golden workbook at its rounding."""
from __future__ import annotations

import csv
import os
import sys

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
FILES = dict(summary="nhws2025_household_summary.csv", persons="nhws2025_persons.csv",
             rw="nhws2025_replicate_weights.csv", crm="crm_clients_20270112.csv", ria="ria_pension_accounts_20261231.csv",
             noms="ria_nominations_20261231.csv", ann="lifetime_annuity_lives_20270112.csv")
CUTS = [35, 45, 55, 65, 75]


def group(age):
    return int(np.searchsorted(CUTS, age, side="right"))


def wmed(values, weights):
    order = np.lexsort((np.arange(len(values)), values))
    v, w = values[order], weights[order]
    cum = np.cumsum(w)
    return v[np.searchsorted(cum, cum[-1] / 2.0, side="left")]


def survey(folder):
    summ = pd.read_csv(os.path.join(folder, FILES["summary"]), dtype={"hhid": str})
    piv = summ.pivot(index="hhid", columns="imp", values="networth").sort_index()
    w = summ[summ.imp == 1].set_index("hhid")["wgt"].reindex(piv.index).to_numpy()
    rw = pd.read_csv(os.path.join(folder, FILES["rw"]), dtype={"hhid": str}).set_index("hhid").reindex(piv.index)
    per = pd.read_csv(os.path.join(folder, FILES["persons"]), dtype={"hhid": str})
    age = per[per.relationship <= 3].groupby("hhid")["age"].max().reindex(piv.index).to_numpy()
    g = np.array([group(a) for a in age])
    nw = piv.to_numpy().T
    R = rw.to_numpy().T
    meds = np.zeros((5, 6))
    repm = np.zeros((R.shape[0], 6))
    for k in range(6):
        sel = g == k
        for m in range(5):
            meds[m, k] = wmed(nw[m][sel], w[sel])
        for r in range(R.shape[0]):
            repm[r, k] = wmed(nw[0][sel], R[r][sel])
    out = []
    for k in range(6):
        u = np.mean((repm[:, k] - meds[0, k]) ** 2)
        b = np.var(meds[:, k], ddof=1)
        out.append(dict(median=meds[:, k].mean(), se=np.sqrt(u + 1.2 * b)))
    pairs = []
    for k in range(5):
        dm = meds[:, k + 1] - meds[:, k]
        dr = repm[:, k + 1] - repm[:, k]
        se = np.sqrt(np.mean((dr - dm[0]) ** 2) + 1.2 * np.var(dm, ddof=1))
        pairs.append(dict(diff=dm.mean(), se=se, apart=abs(dm.mean() / se) > 1.96))
    return out, pairs


def book(folder):
    asat = pd.Timestamp("2026-12-31")
    crm = pd.read_csv(os.path.join(folder, FILES["crm"]), dtype=str)
    crm["born"] = pd.to_datetime(crm["date_of_birth"]).dt.year
    crm["dod"] = pd.to_datetime(crm["date_of_death"])
    crm["alive"] = crm["dod"].isna() | (crm["dod"] > asat)
    merged = dict(zip(crm.client_id, crm.merged_into))
    real = crm[crm.status != "Merged"]
    hh_of = dict(zip(real.client_id, real.household_id))
    born = dict(zip(crm.client_id, crm.born))
    alive = dict(zip(crm.client_id, crm.alive))
    hh_members = real.groupby("household_id")["client_id"].apply(list).to_dict()
    ria = pd.read_csv(os.path.join(folder, FILES["ria"]), dtype=str)
    ria["member"] = [merged[m] if isinstance(merged.get(m), str) else m for m in ria.member_id]
    noms = pd.read_csv(os.path.join(folder, FILES["noms"]), dtype=str)
    noms["nborn"] = pd.to_datetime(noms["nominee_date_of_birth"]).dt.year
    partner_noms = noms[noms.relationship.isin(["Spouse", "De facto partner"])]
    rev = noms[noms.nomination_type == "Reversionary"].set_index("account_no")

    def known_partner(cid):
        ms = hh_members.get(hh_of[cid], [cid])
        if len(ms) > 1:
            return None
        accs = ria.loc[ria.member == cid, "account_no"]
        hit = partner_noms[partner_noms.account_no.isin(accs)]
        return None if hit.empty else int(hit.nborn.iloc[0])

    def placement(cid, extra=None):
        years = [born[m] for m in hh_members.get(hh_of[cid], [cid]) if alive[m]]
        kp = known_partner(cid) if extra is None else extra
        if kp is not None:
            years.append(kp)
        return group(2026 - min(years))

    A = {k: [set(), 0.0, 0.0] for k in range(6)}
    for row in ria.itertuples():
        cid = row.member
        if alive[cid]:
            k, key = placement(cid), hh_of[cid]
        elif row.account_no in rev.index:
            others = [m for m in hh_members.get(hh_of[cid], []) if m != cid and alive[m]]
            if others:
                k, key = group(2026 - min(born[m] for m in others)), hh_of[cid]
            else:
                k, key = group(2026 - int(rev.at[row.account_no, "nborn"])), "rev:" + row.account_no
        else:
            continue
        A[k][0].add(key)
        A[k][1] += float(row.balance_31_12_2026)
        A[k][2] += float(row.annual_payment_2026_27)
    ann = pd.read_csv(os.path.join(folder, FILES["ann"]), dtype=str)
    for c in ("commenced", "payments_start", "date_of_death"):
        ann[c] = pd.to_datetime(ann[c])
    ann["annual_amount"] = ann["annual_amount"].astype(float)
    ann["alive"] = ann["date_of_death"].isna() | (ann["date_of_death"] > asat)
    ann = ann[(ann.commenced <= asat) & (ann.payments_start <= asat)]
    rate = {"CPI": 1.031, "Fixed 2.5%": 1.025, "Level": 1.0}
    Bk = {k: [set(), 0.0] for k in range(6)}
    for no, grp in ann.groupby("contract_no"):
        grp = grp.sort_values("life_no")
        first = grp.iloc[0]
        second = grp.iloc[1] if len(grp) == 2 else None
        if first.alive:
            payee, partner = first, (second if second is not None and second.alive else None)
        elif second is not None and second.alive:
            payee, partner = second, None
        else:
            continue
        amount = payee.annual_amount
        start = payee.commenced
        if start.month == 1 and start.day <= 12:
            anniv = pd.Timestamp(2027, 1, start.day)
            if payee.life_no == "2" or pd.isna(payee.date_of_death) or payee.date_of_death > anniv:
                amount /= rate[payee.indexation]
        extra = int(pd.Timestamp(partner.date_of_birth).year) if partner is not None else None
        k = placement(payee.client_id, extra)
        Bk[k][0].add(hh_of[payee.client_id])
        Bk[k][1] += amount
    return A, Bk


def main(target, golden):
    import openpyxl
    groups, pairs = survey(target)
    A, Bk = book(target)
    wb = openpyxl.load_workbook(os.path.join(golden, "retirement_segments_2027.xlsx"), data_only=True)
    ag = list(wb["Age groups"].iter_rows(min_row=2, values_only=True))
    pr = list(wb["Neighbouring pairs"].iter_rows(min_row=2, values_only=True))
    bk = list(wb["Our retirement book"].iter_rows(min_row=2, values_only=True))
    th = lambda x: int(np.floor(x / 1000 + 0.5)) * 1000
    n = 0
    for k in range(6):
        assert ag[k][2] == th(groups[k]["median"]) and ag[k][3] == th(groups[k]["se"]), (k, ag[k], groups[k])
        n += 2
    for k in range(5):
        assert pr[k][2] == th(pairs[k]["diff"]) and pr[k][3] == th(pairs[k]["se"]), (k, pr[k], pairs[k])
        assert (pr[k][7] == "Kept apart") == pairs[k]["apart"], (k, pr[k])
        n += 3
    for i, k in enumerate([3, 4, 5]):
        row = bk[i]
        assert row[1] == len(A[k][0]) and row[2] == th(A[k][1]) and row[3] == th(A[k][2]), (k, row, A[k])
        assert row[4] == len(Bk[k][0]) and row[5] == th(Bk[k][1]), (k, row, Bk[k])
        n += 5
    segs = sum(1 for p in pairs if p["apart"]) + 1
    assert segs == 5 and [p["apart"] for p in pairs] == [True, True, True, True, False], pairs
    print("verify: %d figures recomputed from the bundle on an independent path, every one equal to the golden "
          "workbook; %d segments" % (n + 1, segs))


if __name__ == "__main__":
    task = os.path.dirname(HERE)
    main(os.path.join(task, "target"), os.path.join(task, "golden"))
