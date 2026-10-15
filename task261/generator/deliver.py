"""The golden, computed from the shipped pack alone: each age group's median household net worth (households placed by
the older of the reference person and partner), each neighbouring difference with its standard error taken from the
difference itself on the replicate weights and implicates (user guide 7), the segments the committee's rule gives, and
the asks on Kelvane Life's retirement income book; then the committee paper, the workbook and the chart."""
from __future__ import annotations

import datetime as dt
import os

import numpy as np
import pandas as pd

import common as K
import est as E

AS_AT = pd.Timestamp(K.AS_AT)
JAN_RATE = {"CPI": 0.031, "Fixed 2.5%": 0.025, "Level": 0.0}     # the key facts sheet's increases, Nov 2026 to Jan 2027


def load_survey(folder):
    S = pd.read_csv(os.path.join(folder, K.SUMMARY), dtype={"hhid": str})
    P = pd.read_csv(os.path.join(folder, K.PERSONS), dtype={"hhid": str})
    RW = pd.read_csv(os.path.join(folder, K.REPW), dtype={"hhid": str}).set_index("hhid")
    ids = np.sort(S["hhid"].unique())
    nw = np.stack([S[S["imp"] == m].set_index("hhid").loc[ids, "networth"].to_numpy(float)
                   for m in range(1, K.IMPLICATES + 1)])
    w = S[S["imp"] == 1].set_index("hhid").loc[ids, "wgt"].to_numpy(float)
    R = RW.loc[ids].to_numpy(float).T
    couple_age = P[P["relationship"].isin([1, 2, 3])].groupby("hhid")["age"].max()
    ref_age = P[P["relationship"] == 1].set_index("hhid")["age"]
    older = couple_age.reindex(ids).to_numpy()
    return dict(ids=ids, nw=nw, w=w, R=R, older=older, ref=ref_age.reindex(ids).to_numpy(), S=S, P=P)


def main_call(sv, band):
    """Medians per group and the neighbouring differences, each difference's error from the difference itself."""
    groups = []
    per = []
    for g in range(6):
        est, rep = sv.median(band == g)
        th, u, b, t, se = sv.combine(est, rep)
        groups.append(dict(name=K.BAND_NAMES[g], median=th, se=se, est=est, rep=rep,
                           households=float(sv.w[band == g].sum())))
        per.append((est, rep))
    pairs = []
    for g in range(5):
        (e0, r0), (e1, r1) = per[g], per[g + 1]
        de, dr = e1 - e0, r1 - r0
        th, u, b, t, se = E.Survey.combine(de, dr)
        quad = np.sqrt(groups[g]["se"] ** 2 + groups[g + 1]["se"] ** 2)
        corr = float(np.corrcoef(r0, r1)[0, 1])
        pairs.append(dict(lo=K.BAND_NAMES[g], hi=K.BAND_NAMES[g + 1], diff=th, se=se, ratio=th / se,
                          apart=abs(th / se) > 1.96, ci=(th - 1.96 * se, th + 1.96 * se), quad=quad, corr=corr))
    segs, cur = [], [0]
    for g in range(5):
        if pairs[g]["apart"]:
            segs.append(cur)
            cur = [g + 1]
        else:
            cur.append(g + 1)
    segs.append(cur)
    return groups, pairs, segs


def seg_name(seg):
    lo, hi = K.BANDS[seg[0]][0], K.BANDS[seg[-1]][1]
    if seg[0] == 0:
        return "Under %d" % (K.BANDS[seg[-1]][1] + 1)
    if hi >= 120:
        return "%d and over" % lo
    return "%d to %d" % (lo, hi)


# ---------------------------------------------------------------------------------------------- the asks
def load_book(folder):
    crm = pd.read_csv(os.path.join(folder, K.CRM), parse_dates=["date_of_birth", "date_of_death"], dtype=str)
    crm["date_of_birth"] = pd.to_datetime(crm["date_of_birth"])
    crm["date_of_death"] = pd.to_datetime(crm["date_of_death"])
    ria = pd.read_csv(os.path.join(folder, K.RIA), dtype={"account_no": str, "member_id": str})
    noms = pd.read_csv(os.path.join(folder, K.NOMS), dtype={"account_no": str}, parse_dates=["nominee_date_of_birth"])
    ann = pd.read_csv(os.path.join(folder, K.ANN), dtype={"contract_no": str, "client_id": str},
                      parse_dates=["date_of_birth", "date_of_death", "commenced", "payments_start"])
    return crm, ria, noms, ann


class Book:
    """Kelvane's clients as the register, the pension extract, the nominations and the annuity listing hold them, with
    each household's members and older partner at 31 December 2026."""

    def __init__(self, crm, ria, noms):
        self.c = crm.set_index("client_id")
        self.merged = self.c["merged_into"].dropna()
        live = crm[crm["status"] != "Merged"]
        self.members = live.groupby("household_id")["client_id"].apply(list).to_dict()
        self.ria = ria
        self.rev = noms[noms["nomination_type"] == "Reversionary"].set_index("account_no")
        self.spouse = noms[noms["relationship"].isin(["Spouse", "De facto partner"])]
        self.owner = {a: self.resolve(m) for a, m in zip(ria["account_no"], ria["member_id"])}

    def resolve(self, cid):
        return self.merged.get(cid, cid)

    def alive(self, cid):
        d = self.c.at[cid, "date_of_death"]
        return pd.isna(d) or d > AS_AT

    def age(self, cid):
        return K.AS_AT.year - self.c.at[cid, "date_of_birth"].year

    def household(self, cid):
        return self.c.at[cid, "household_id"], self.members.get(self.c.at[cid, "household_id"], [cid])

    def nominated_partner_age(self, cid):
        """The age of a partner named on a nomination of any of the client's pensions, where the register holds no
        partner for the client."""
        hh, ms = self.household(cid)
        if len(ms) > 1:
            return None
        accs = [a for a, o in self.owner.items() if o == cid]
        sp = self.spouse[self.spouse["account_no"].isin(accs)]
        return None if sp.empty else K.AS_AT.year - sp["nominee_date_of_birth"].iloc[0].year


def ask_a(bk):
    out = {g: dict(hh=set(), bal=0.0, pay=0.0) for g in range(6)}
    for a in bk.ria.itertuples():
        cid = bk.owner[a.account_no]
        if bk.alive(cid):
            hh, ms = bk.household(cid)
            ages = [bk.age(m) for m in ms if bk.alive(m)]
            npa = bk.nominated_partner_age(cid)
            if npa is not None:
                ages.append(npa)
            key = hh
        elif a.account_no in bk.rev.index:
            # the pension continued to the reversionary beneficiary from the date of death (PDS, section 5)
            hh, ms = bk.household(cid)
            living = [m for m in ms if m != cid and bk.alive(m)]
            if living:
                ages, key = [bk.age(m) for m in living], hh
            else:
                ages = [K.AS_AT.year - bk.rev.at[a.account_no, "nominee_date_of_birth"].year]
                key = "reversion " + a.account_no
        else:
            continue                       # the pension stopped at the member's death (PDS, section 5)
        g = int(K.band_of(max(ages)))
        out[g]["hh"].add(key)
        out[g]["bal"] += a.balance_31_12_2026
        out[g]["pay"] += a.annual_payment_2026_27
    return {g: dict(households=len(v["hh"]), balance=v["bal"], payments=v["pay"]) for g, v in out.items()}


def ask_b(bk, ann):
    out = {g: dict(hh=set(), income=0.0) for g in range(6)}
    c = ann[(ann["commenced"] <= AS_AT) & (ann["payments_start"] <= AS_AT)]
    alive31 = c["date_of_death"].isna() | (c["date_of_death"] > AS_AT)
    c = c.assign(alive31=alive31)
    for k, grp in c.groupby("contract_no"):
        grp = grp.sort_values("life_no")
        l1 = grp.iloc[0]
        l2 = grp.iloc[1] if len(grp) > 1 else None
        if l1["alive31"]:
            pay, other = l1, (l2 if l2 is not None and l2["alive31"] else None)
        elif l2 is not None and l2["alive31"]:
            pay, other = l2, None
        else:
            continue
        amt = float(pay["annual_amount"])
        s = pay["commenced"]
        if s.month == 1 and s.day <= K.PULL.day:
            ann_day = pd.Timestamp(2027, 1, s.day)
            # the listing is at 12 January 2027: an increase on a January anniversary before the pull is taken back
            if pay["life_no"] == 2 or pd.isna(pay["date_of_death"]) or pay["date_of_death"] > ann_day:
                amt = amt / (1 + JAN_RATE[pay["indexation"]])
        cid = pay["client_id"]
        hh, ms = bk.household(cid)
        ages = [bk.age(m) for m in ms if bk.alive(m)]
        if other is not None:
            ages.append(K.AS_AT.year - other["date_of_birth"].year)
        else:
            npa = bk.nominated_partner_age(cid)
            if npa is not None:
                ages.append(npa)
        g = int(K.band_of(max(ages)))
        out[g]["hh"].add(hh)
        out[g]["income"] += amt
    return {g: dict(households=len(v["hh"]), income=v["income"]) for g, v in out.items()}


def thousands(x):
    return int(np.floor(x / 1000 + 0.5))


def compute(folder):
    s = load_survey(folder)
    sv = E.Survey(s["nw"], s["w"], s["R"])
    band = K.band_of(s["older"])
    groups, pairs, segs = main_call(sv, band)
    crm, ria, noms, ann = load_book(folder)
    bk = Book(crm, ria, noms)
    A = ask_a(bk)
    Bk = ask_b(bk, ann)
    names = [seg_name(x) for x in segs]
    fig = {"call": "Five segments: " + ", ".join(names) if len(segs) == 5 else "%d segments" % len(segs),
           "segments": names}
    for g in range(6):
        fig["median_%d" % g] = k_(groups[g]["median"])
        fig["se_%d" % g] = k_(groups[g]["se"])
    for g in range(5):
        p = pairs[g]
        fig["diff_%d" % g] = k_(p["diff"])
        fig["se_diff_%d" % g] = k_(p["se"])
        fig["ratio_%d" % g] = "%.1f" % p["ratio"]
        fig["ci_%d" % g] = "%s to %s" % (k_(p["ci"][0]), k_(p["ci"][1]))
        fig["apart_%d" % g] = "kept apart" if p["apart"] else "one segment"
    fig["se_quadrature_3"] = k_(pairs[3]["quad"])
    fig["ratio_quadrature_3"] = "%.1f" % (pairs[3]["diff"] / pairs[3]["quad"])
    fig["corr_3"] = "%.2f" % pairs[3]["corr"]
    fig["diff_abs_4"] = k_(abs(pairs[4]["diff"]))
    fig["ratio_abs_4"] = "%.1f" % abs(pairs[4]["ratio"])
    fig["constants"] = dict(threshold="1.96", replicates="200", areas="120", implicate_factor="1.2", interval="95",
                            age_bounds="35 44 45 54 55 64 65 74 75")
    for g in K.OLDER:
        fig["A_households_%d" % g] = format(A[g]["households"], ",")
        fig["A_balance_%d" % g] = k_(A[g]["balance"])
        fig["A_payments_%d" % g] = k_(A[g]["payments"])
        fig["B_households_%d" % g] = format(Bk[g]["households"], ",")
        fig["B_income_%d" % g] = k_(Bk[g]["income"])
    return dict(groups=groups, pairs=pairs, segs=segs, names=names, A=A, B=Bk, figures=fig)


# ---------------------------------------------------------------------------------------------- the deliverables
def k_(x):
    v = thousands(x) * 1000
    return ("-$" if v < 0 else "$") + format(abs(v), ",")


def write_paper(path, res):
    import docx
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
    G, Pp, A, Bk = res["groups"], res["pairs"], res["A"], res["B"]
    doc = docx.Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    sec = doc.sections[0]
    sec.header.paragraphs[0].text = "Kelvane Life | Retirement Product Committee, %s | Paper RPC-2027-03" % (
        K.COMMITTEE.strftime("%-d %B %Y"))
    sec.header.paragraphs[0].runs[0].font.size = Pt(8)
    sec.header.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x5A, 0x64, 0x6D)
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = fp.add_run("Page ")
    run.font.size = Pt(8)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        el = OxmlElement("w:fldChar" if kind else "w:instrText")
        if kind:
            el.set(qn("w:fldCharType"), kind)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = text
        r2 = fp.add_run()
        r2.font.size = Pt(8)
        r2._r.append(el)
    t = doc.add_paragraph()
    r = t.add_run("Five segments for the 2027 retirement range: 55 to 64 stands on its own; 65 and over is one segment")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0x0F, 0x4C, 0x5C)
    meta = doc.add_paragraph("For decision. From Insights and Analytics to %s, Head of Retirement Products, and the "
                             "Retirement Product Committee (chair %s). %s." % (
                                 K.PEOPLE["head"], K.PEOPLE["chair"], dt.date(2027, 1, 20).strftime("%-d %B %Y")))
    meta.runs[0].font.size = Pt(8.5)
    meta.runs[0].font.color.rgb = RGBColor(0x5A, 0x64, 0x6D)

    def h(text):
        p = doc.add_paragraph()
        rr = p.add_run(text)
        rr.bold = True
        rr.font.size = Pt(11)
        p.paragraph_format.space_before = Pt(8)

    h("Recommendation")
    doc.add_paragraph(
        "Set the 2027 range on five segments: %s. Under the committee's rule every boundary but one holds, and 65 to 74 "
        "shares a segment with 75 and over." % ", ".join(res["names"]))
    d3 = Pp[3]
    d4 = Pp[4]
    doc.add_paragraph(
        "The boundary that decides the structure is the one at 65. Households whose older partner is 55 to 64 have a "
        "median net worth of %s against %s for 65 to 74, a difference of %s whose standard error is %s, so the "
        "difference is %.1f times its standard error and the two groups stay apart. Between 65 to 74 and 75 and over "
        "the medians differ by %s against a standard error of %s (%.1f times), so they share a segment." % (
            k_(G[3]["median"]), k_(G[4]["median"]), k_(d3["diff"]), k_(d3["se"]), d3["ratio"], k_(abs(d4["diff"])),
            k_(d4["se"]), abs(d4["ratio"])))
    doc.add_paragraph(
        "Each difference's standard error is estimated on the difference itself: recomputed with each of the survey's "
        "200 replicate weights, with the five implicates' variance added, as the user guide sets out for any estimate. "
        "The age groups are drawn from the same 120 survey areas, and where an area's housing is worth more both groups' "
        "medians rise together: across the replicates the 55 to 64 and 65 to 74 medians move together (correlation "
        "%.2f), so their difference varies much less than the two groups' own standard errors suggest. Added together "
        "as if the groups were separate samples, those errors would put the 65 boundary's standard error at %s, merge 55 "
        "to 64 with 65 and over and leave the range with four segments." % (d3["corr"], k_(d3["quad"])))
    h("The evidence for each boundary")
    tb = doc.add_table(rows=1, cols=6)
    tb.style = "Light List Accent 1"
    tb.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, txt in enumerate(["Neighbouring groups", "Difference in median", "Standard error", "Times its error",
                             "95% interval", "Committee's rule"]):
        tb.rows[0].cells[i].text = txt
    for p in Pp:
        row = tb.add_row().cells
        vals = ["%s / %s" % (p["lo"], p["hi"]), k_(p["diff"]), k_(p["se"]), "%.1f" % p["ratio"],
                "%s to %s" % (k_(p["ci"][0]), k_(p["ci"][1])), "Kept apart" if p["apart"] else "One segment"]
        for i, v in enumerate(vals):
            row[i].text = v
            if i in (1, 2, 3):
                row[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph()
    tb = doc.add_table(rows=1, cols=4)
    tb.style = "Light List Accent 1"
    for i, txt in enumerate(["Age group (older partner)", "Households (weighted)", "Median net worth",
                             "Standard error"]):
        tb.rows[0].cells[i].text = txt
    for g in G:
        row = tb.add_row().cells
        for i, v in enumerate([g["name"], format(int(round(g["households"], -3)), ","), k_(g["median"]), k_(g["se"])]):
            row[i].text = v
            if i:
                row[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    note = doc.add_paragraph("Source: National Household Wealth Survey 2025, public-use release (Household Finance "
                             "Research Centre, September 2026); medians weighted, averaged over five implicates.")
    note.runs[0].font.size = Pt(8)
    h("Our retirement income customers in the three groups from 55")
    doc.add_paragraph(
        "Placed the same way, by the older partner's age at 31 December 2026, our Retirement Income Account and lifetime "
        "annuity customers sit as below. A pension whose member died before 31 December with a reversionary "
        "beneficiary nominated is counted with the survivor, to whom it continued; a partner who is not our client is "
        "placed from the date of birth on the member's nomination.")
    tb = doc.add_table(rows=1, cols=6)
    tb.style = "Light List Accent 1"
    for i, txt in enumerate(["Age group", "Households with a pension", "Pension balances", "Payments set for 2026-27",
                             "Households with an annuity in payment", "Annuity income a year"]):
        tb.rows[0].cells[i].text = txt
    for g in K.OLDER:
        row = tb.add_row().cells
        vals = [K.BAND_NAMES[g], format(A[g]["households"], ","), k_(A[g]["balance"]), k_(A[g]["payments"]),
                format(Bk[g]["households"], ","), k_(Bk[g]["income"])]
        for i, v in enumerate(vals):
            row[i].text = v
            if i:
                row[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    note = doc.add_paragraph("Sources: client register and annuity valuation listing (pulled 12 January 2027), "
                             "Retirement Income Account extract and nominations (31 December 2026). Annuity income is the "
                             "annual rate being paid at 31 December 2026, before the January anniversary increases; "
                             "deferred annuities not yet paying are left out.")
    note.runs[0].font.size = Pt(8)
    for p in doc.paragraphs:
        p.paragraph_format.space_after = Pt(5)
    doc.core_properties.author = "Insights and Analytics"
    doc.save(path)


def write_xlsx(path, res):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    G, Pp, A, Bk = res["groups"], res["pairs"], res["A"], res["B"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Structure"
    ws.append(["2027 retirement range: segments"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Segment", "Age groups (older partner)"])
    for i, (nm, sg) in enumerate(zip(res["names"], res["segs"]), 1):
        ws.append(["Segment %d: %s" % (i, nm), ", ".join(K.BAND_NAMES[g] for g in sg)])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 40
    for c in ws[2]:
        c.font = Font(bold=True)
    s = wb.create_sheet("Age groups")
    s.append(["Age group (older partner)", "Households", "Median net worth ($)", "Standard error ($)"])
    for g in G:
        s.append([g["name"], int(round(g["households"], -3)), thousands(g["median"]) * 1000, thousands(g["se"]) * 1000])
    s = wb.create_sheet("Neighbouring pairs")
    s.append(["Younger group", "Older group", "Difference in median ($)", "Standard error of the difference ($)",
              "Difference / standard error", "95% interval, low ($)", "95% interval, high ($)", "Committee's rule"])
    for p in Pp:
        s.append([p["lo"], p["hi"], thousands(p["diff"]) * 1000, thousands(p["se"]) * 1000, round(p["ratio"], 1),
                  thousands(p["ci"][0]) * 1000, thousands(p["ci"][1]) * 1000, "Kept apart" if p["apart"] else "One segment"])
    s = wb.create_sheet("Our retirement book")
    s.append(["Age group (older partner, 31 December 2026)", "Households with a Retirement Income Account pension",
              "Pension balances ($)", "Annual payments set for 2026-27 ($)", "Households with a lifetime annuity in payment",
              "Lifetime annuity income a year ($)"])
    for g in K.OLDER:
        s.append([K.BAND_NAMES[g], A[g]["households"], thousands(A[g]["balance"]) * 1000,
                  thousands(A[g]["payments"]) * 1000, Bk[g]["households"], thousands(Bk[g]["income"]) * 1000])
    s.append(["Three groups from 55", sum(A[g]["households"] for g in K.OLDER),
              thousands(sum(A[g]["balance"] for g in K.OLDER)) * 1000, thousands(sum(A[g]["payments"] for g in K.OLDER)) * 1000,
              sum(Bk[g]["households"] for g in K.OLDER), thousands(sum(Bk[g]["income"] for g in K.OLDER)) * 1000])
    for c in s[s.max_row]:
        c.font = Font(bold=True)
    s = wb.create_sheet("Notes")
    notes = ["Survey: National Household Wealth Survey 2025, public-use release (Household Finance Research Centre, "
             "September 2026).",
             "Households are placed by the age of the older of the reference person and partner (the brief).",
             "Medians are weighted and averaged over the five implicates; standard errors use the 200 replicate weights "
             "on implicate 1 with 1.2 times the variance between implicates added (user guide, section 7).",
             "Each difference's standard error is taken from the difference itself, recomputed on every replicate "
             "weight and implicate. Committee's rule: neighbouring groups are kept apart where the difference is more "
             "than 1.96 times its standard error.",
             "Our book: client register and annuity listing pulled 12 January 2027; pension extract and nominations at "
             "31 December 2026. A pension whose member died before 31 December with a reversionary nominated is counted "
             "with the survivor; annuity income is the annual rate at 31 December 2026, deferred annuities not yet "
             "paying left out.",
             "Dollar figures rounded to the nearest thousand."]
    for t in notes:
        s.append([t])
    s.column_dimensions["A"].width = 130
    for sh in wb.worksheets[1:4]:
        for c in sh[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.fill = PatternFill("solid", fgColor="DCE6EC")
        sh.freeze_panes = "A2"
        for col in "ABCDEFGH":
            sh.column_dimensions[col].width = 19
        sh.column_dimensions["A"].width = 30
        for row in sh.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, int) and c.value >= 1000:
                    c.number_format = "#,##0"
                elif isinstance(c.value, float):
                    c.number_format = "0.0"
    wb.save(path)


def write_chart(path, res):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams["svg.hashsalt"] = "task261"
    plt.rcParams["font.family"] = "DejaVu Sans"
    Pp = res["pairs"]
    fig, ax = plt.subplots(figsize=(8.0, 4.5), dpi=100)
    ys = np.arange(len(Pp))[::-1]
    for y, p in zip(ys, Pp):
        col = "#0F4C5C" if p["apart"] else "#B5532F"
        ax.errorbar(p["diff"] / 1000, y, xerr=[[(p["diff"] - p["ci"][0]) / 1000], [(p["ci"][1] - p["diff"]) / 1000]],
                    fmt="o", color=col, ecolor=col, capsize=4, lw=1.8, ms=6)
        lab = "%s%s,000" % ("+$" if p["diff"] > 0 else "-$", format(abs(thousands(p["diff"])), ","))
        ax.text(p["ci"][1] / 1000 + 14, y, lab + ("  kept apart" if p["apart"] else "  one segment"), va="center",
                fontsize=8.3, color=col)
    ax.axvline(0, color="#333333", lw=1)
    ax.set_yticks(ys)
    ax.set_yticklabels(["%s / %s" % (p["lo"], p["hi"]) for p in Pp], fontsize=9)
    ax.set_xlim(-600, 760)
    ax.set_xlabel("Difference in median household net worth, older group less younger ($'000)", fontsize=8.6)
    ax.set_title("Four of the five boundaries clear zero: the range has five segments", fontsize=10.5, loc="left")
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="x", color="#DDDDDD", lw=0.6)
    fig.text(0.012, 0.015, "Bars: 95 per cent interval, the difference plus or minus 1.96 standard errors of the difference "
             "(survey replicate weights and implicates).\nHouseholds placed by the older partner. Source: National "
             "Household Wealth Survey 2025, public-use release.", fontsize=7.0, color="#5A646D", va="bottom")
    fig.tight_layout(rect=(0, 0.07, 1, 1))
    fig.savefig(path, format="svg", metadata={"Date": None, "Creator": None})
    plt.close(fig)
