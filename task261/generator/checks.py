"""task261's generator checks: the input gates, the bundle against the generator's clean truth, the three rules
executed (R1 to R3c), every semantic fork row (SF1 to SF16), the ladder (every error construction and placement, its
segments, the margins), the asks' stops and cracker sheet, the separation line, the distractors, the goldens against
figures.json, and the realism sweep (H1, H21, dates, round figures, personas)."""
from __future__ import annotations

import json
import os
import re
import shutil
import tempfile
import zipfile

import numpy as np
import pandas as pd

import common as K
import deliver as G
import est as E
import insurer as I
import stops as T
import tune as TU

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))


class Checks:
    def __init__(self):
        self.n = 0

    def __call__(self, label, cond, detail=""):
        if not cond:
            raise AssertionError("%s: %s" % (label, detail))
        self.n += 1


def pdf_text(path):
    from pypdf import PdfReader
    return "\n".join(p.extract_text() or "" for p in PdfReader(path).pages)


def docx_text(path):
    import docx
    d = docx.Document(path)
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            out.append(" | ".join(c.text for c in r.cells))
    for s in d.sections:
        out += [p.text for p in s.header.paragraphs]
    return "\n".join(out)


def xlsx_text(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            out.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def segments(z):
    return tuple(abs(v) > 1.96 for v in z)


GOLD_SPLITS = (True, True, True, True, False)
DEFAULT_SPLITS = (True, True, True, False, False)


def arrays(s):
    return dict(nw=s["nw"], wgt=s["w"], repw=s["R"], older=s["older"], ref=s["ref"])


def zc(nw, w, R, band, rep_all=False, interp=False, b_factor=1.2):
    """Each neighbouring difference over the error of the difference itself (the decisive construction), with
    convention switches for the fork table."""
    def med(x, ww):
        if not interp:
            return E.wmedian(x, ww)
        o = np.argsort(x, kind="stable")
        xs = x[o]
        out = []
        for row in np.atleast_2d(ww):
            c = np.cumsum(row[o])
            p = (c - row[o] / 2) / c[-1]
            out.append(np.interp(0.5, p, xs))
        return np.array(out)
    est = np.array([[med(nw[m][band == g], w[band == g])[0] for g in range(6)] for m in range(5)])
    reps = [np.stack([med(nw[m][band == g], R[:, band == g]) for g in range(6)], 1) for m in (range(5) if rep_all else [0])]
    z, se = [], []
    for g in range(5):
        d = est[:, g + 1] - est[:, g]
        u = np.mean([(((rp[:, g + 1] - rp[:, g]) - d[m]) ** 2).mean() for m, rp in zip(range(5) if rep_all else [0], reps)])
        v = u + b_factor * d.var(ddof=1)
        z.append(d.mean() / np.sqrt(v))
        se.append(np.sqrt(v))
    return np.array(z), est.mean(0), np.array(se)


def run(target, golden, st, meta_path):
    ck = Checks()
    d, B = st["S"]["d"], st["S"]["B"]
    res = st["res"]
    fig = res["figures"]
    files = sorted(os.listdir(target))

    # ------------------------------------------------------------------ input gates
    ck("gate: 10 or more files", len(files) >= 10, len(files))
    ck("gate: 3 or more formats", len({f.rsplit(".", 1)[1] for f in files}) >= 3, files)
    rows = {f: sum(1 for _ in open(os.path.join(target, f), encoding="utf-8")) - 1 for f in files if f.endswith(".csv")}
    ck("gate: one file of 25,000 rows or more", max(rows.values()) >= 25000, rows)
    meta = json.load(open(meta_path))
    ck("gate: two or more distractors declared", len(meta["distractor_files"]) >= 2, meta)
    ck("gate: every declared file ships", all(f in files for f in meta["distractor_files"] + meta["call_files"] +
                                              meta["ask_files"]), meta)
    ck("gate: every shipped file declared", set(files) == set(meta["distractor_files"] + meta["call_files"] +
                                                             meta["ask_files"]), set(files))
    ck("gate: no file name broadcasts the device or the stump", not any(re.search(
        r"distractor|decoy|answer|golden|issue|error|covarian|contrast|revers|merge|dup", f, re.I) for f in files), files)

    # ------------------------------------------------------------------ the bundle against the clean truth
    s = G.load_survey(target)
    order = np.argsort(d["hhid"])
    ck("bundle: survey households", list(s["ids"]) == list(np.array(d["hhid"])[order]), len(s["ids"]))
    ck("bundle: net worth on every implicate", np.array_equal(s["nw"], d["nw"][:, order]), "")
    ck("bundle: weights", np.allclose(s["w"], d["wgt"][order]), "")
    ck("bundle: replicate weights", np.allclose(s["R"], d["repw"][:, order]), "")
    older_true = np.minimum(d["older"][order], 85)
    ck("bundle: older partner's age from the person file", np.array_equal(s["older"], older_true), "")
    A_t, Bn_t, Bh_t = I.golden(B)
    for g in K.OLDER:
        ck("bundle: Ask A households %s" % K.BAND_SHORT[g], res["A"][g]["households"] == len(A_t[g]["hh"]),
           (res["A"][g]["households"], len(A_t[g]["hh"])))
        ck("bundle: Ask A balance %s" % K.BAND_SHORT[g], abs(res["A"][g]["balance"] - A_t[g]["bal"]) < 0.01, "")
        ck("bundle: Ask A payments %s" % K.BAND_SHORT[g], abs(res["A"][g]["payments"] - A_t[g]["pay"]) < 0.01, "")
        ck("bundle: Ask B households %s" % K.BAND_SHORT[g], res["B"][g]["households"] == len(Bh_t[g]),
           (res["B"][g]["households"], len(Bh_t[g])))
        ck("bundle: Ask B income %s (deflation residual under a dollar)" % K.BAND_SHORT[g],
           abs(res["B"][g]["income"] - Bn_t[g]) < 1.0, (res["B"][g]["income"], Bn_t[g]))

    # ------------------------------------------------------------------ R1: the field's raw material ships
    for f in (K.SUMMARY, K.PERSONS, K.REPW):
        head = open(os.path.join(target, f), encoding="utf-8").readline().lower()
        ck("R1: %s holds microdata, no estimate" % f,
           not re.search(r"median|stderr|se_|rse|diff|ratio|segment|z_", head), head[:120])
    ck("R1: the summary is household by implicate", rows[K.SUMMARY] == 5 * len(s["ids"]), rows[K.SUMMARY])
    ck("R1: 200 replicate weights per household", s["R"].shape == (K.REPLICATES, len(s["ids"])), s["R"].shape)

    # ------------------------------------------------------------------ the ladder on the bundle
    o = TU.routes(arrays(s))
    band = K.band_of(s["older"])
    ck("ladder R3 (decisive): the difference on the replicates gives five segments", segments(o["older"]["zc"]) == GOLD_SPLITS,
       o["older"]["zc"])
    ck("ladder R0 (default): quadrature gives four segments", segments(o["older"]["zq"]) == DEFAULT_SPLITS, o["older"]["zq"])
    ck("ladder R1: the reference person's age, quadrature: not the golden", segments(o["ref"]["zq"]) != GOLD_SPLITS,
       o["ref"]["zq"])
    ck("ladder R1': the reference person's age, the difference on the replicates: not the golden",
       segments(o["ref"]["zc"]) != GOLD_SPLITS, o["ref"]["zc"])
    ck("ladder R2: errors without the design (Woodruff): six segments", all(segments(o["older"]["zs"])), o["older"]["zs"])
    ck("ladder R2: a bootstrap over households: six segments", all(segments(o["older"]["zb"])), o["older"]["zb"])
    ck("ladder R2: the reference person without the design: not the golden", segments(o["ref"]["zs"]) != GOLD_SPLITS,
       o["ref"]["zs"])
    zg = o["older"]["zc"]
    ck("margin: every golden ratio at least 0.3 from 1.96", min(abs(abs(zg) - 1.96)) >= 0.3, zg)
    ck("margin: the default's 65 boundary at least 0.3 under 1.96", 1.96 - o["older"]["zq"][3] >= 0.3, o["older"]["zq"])
    ck("margin: without the design, 65-74 / 75 and over at least 0.3 past 1.96",
       min(abs(o["older"]["zs"][4]), abs(o["older"]["zb"][4])) - 1.96 >= 0.3, (o["older"]["zs"], o["older"]["zb"]))
    ck("ladder: the golden pipeline agrees with the tuning routes", np.allclose([p["ratio"] for p in res["pairs"]], zg),
       ([p["ratio"] for p in res["pairs"]], zg))
    ck("ladder: the decisive pair's replicate correlation is strong (0.5 or more)", res["pairs"][3]["corr"] >= 0.5,
       res["pairs"][3]["corr"])
    ck("ladder: the quadrature is at least 1.5 times the difference's own error at 65", res["pairs"][3]["quad"] >=
       1.5 * res["pairs"][3]["se"], (res["pairs"][3]["quad"], res["pairs"][3]["se"]))
    ck("ladder: every rung files a different structure", len({segments(o["older"]["zc"]), segments(o["older"]["zq"]),
                                                                 segments(o["older"]["zs"])}) == 3, "")

    # ------------------------------------------------------------------ R3a, R3b on the clean truth
    clean = dict(nw=np.stack([d["nw0"] * K.SCALE * np.where(d["nw0"] > 0, np.asarray(K.TUNE)[K.band_of(d["older"])], 1)
                              * np.where(d["nw0"] > 0, np.exp(_older_area(d)), 1.0)] * 5),
                 wgt=d["wgt"], repw=d["repw"], older=d["older"], ref=d["ref"])
    oc = TU.routes(clean, nboot=60)
    ck("R3a: the easy route (quadrature) on the clean truth files four segments", segments(oc["older"]["zq"]) ==
       DEFAULT_SPLITS, oc["older"]["zq"])
    ck("R3b: the decisive step alone on the clean truth files five", segments(oc["older"]["zc"]) == GOLD_SPLITS,
       oc["older"]["zc"])
    ck("R3b: correcting the design alone (from SRS to quadrature) leaves a wrong answer", segments(oc["older"]["zq"]) !=
       GOLD_SPLITS and segments(oc["older"]["zs"]) != GOLD_SPLITS, (oc["older"]["zq"], oc["older"]["zs"]))

    # ------------------------------------------------------------------ R2: every Layer 1 item neutral
    z1, _, _ = zc(s["nw"][[0] * 5], s["w"], s["R"], band)
    ck("R2: implicate 1 alone leaves the call", segments(z1) == GOLD_SPLITS, z1)
    stack_nw = s["nw"].reshape(-1)
    stack_w = np.tile(s["w"], 5) / 5
    stack_b = np.tile(band, 5)
    med_stack = [E.wmedian(stack_nw[stack_b == g], stack_w[stack_b == g])[0] for g in range(6)]
    med_gold = [g["median"] for g in res["groups"]]
    ck("R2: the five implicates stacked at a fifth of the weight move no median by 2 per cent",
       max(abs(a / b - 1) for a, b in zip(med_stack, med_gold)) < 0.02, (med_stack, med_gold))
    keep = s["nw"].min(0) >= 0
    zneg, _, _ = zc(s["nw"][:, keep], s["w"][keep], s["R"][:, keep], band[keep])
    ck("R2: negative net worth dropped leaves the call", segments(zneg) == GOLD_SPLITS, zneg)
    ck("R2: the age top-code at 85 moves no household's group", np.array_equal(K.band_of(s["older"]),
                                                                              K.band_of(d["older"][order])), "")
    ck("R2: replicate weights joined on hhid cover every household", set(pd.read_csv(os.path.join(target, K.REPW),
                                                                                    usecols=["hhid"], dtype=str)["hhid"])
       == set(s["ids"]), "")

    # ------------------------------------------------------------------ R3c: no plant
    texts = {f: _text(os.path.join(target, f)) for f in files if not f.endswith(".csv")}
    joined = "\n".join(texts.values()).lower()
    for w in ("covarian", "correlat", "quadrature", "square root of the sum", "independent", "difference between two",
              "contrast", "five segments", "four segments", "three segments", "six segments"):
        ck("R3c: no shipped document carries '%s'" % w, w not in joined, w)
    ck("R3c: the brief files the rule in standard errors of the difference, no method",
       "1.96 times the standard error of that difference" in texts[K.BRIEF], "")
    for f in (K.SUMMARY, K.REPW):
        lines = open(os.path.join(target, f), encoding="utf-8").read().splitlines()[1:]
        body = [ln.split(",", 1)[1] for ln in lines]
        ck("R3c: no copied block in %s" % f, len(set(body)) > 0.99 * len(body), "")

    # ------------------------------------------------------------------ the semantic fork table
    ck("SF1: the brief places by the older of the reference person and partner",
       "older of its reference person and partner" in texts[K.BRIEF], "")
    ref_b = K.band_of(s["ref"])
    zr, _, _ = zc(s["nw"], s["w"], s["R"], ref_b)
    ck("SF1: the reference person's age files another structure", segments(zr) != GOLD_SPLITS, zr)
    P = s["P"]
    m2 = P[P["relationship"].isin([1, 2])].groupby("hhid")["age"].max().reindex(s["ids"]).to_numpy()
    zm2, _, _ = zc(s["nw"], s["w"], s["R"], K.band_of(m2))
    ck("SF2: married partners only: same call", segments(zm2) == GOLD_SPLITS, zm2)
    ck("SF2: the user guide lists codes 2 and 3 as partners", "2 husband, wife or registered partner; 3 de facto partner"
       in texts[K.GUIDE].replace("\n", " "), "")
    ck("SF3: the user guide defines net worth as assets less debts", "networth is assets less debts" in
       texts[K.GUIDE].replace("\n", " "), "")
    ck("SF3: negatives kept or dropped: same call", segments(zneg) == GOLD_SPLITS, zneg)
    ck("SF4: the user guide requires weights", "Every estimate must be weighted" in texts[K.GUIDE], "")
    zi, _, _ = zc(s["nw"], s["w"], s["R"], band, interp=True)
    ck("SF5: the interpolated median: same call", segments(zi) == GOLD_SPLITS, zi)
    ck("SF6: implicate 1 alone: same call", segments(z1) == GOLD_SPLITS, z1)
    za, _, _ = zc(s["nw"], s["w"], s["R"], band, rep_all=True)
    ck("SF7: replicates on all five implicates: same call", segments(za) == GOLD_SPLITS, za)
    zb0, _, _ = zc(s["nw"], s["w"], s["R"], band, b_factor=0.0)
    ck("SF8: the between-implicate variance left out: same call", segments(zb0) == GOLD_SPLITS, zb0)
    allz = np.concatenate([zg, z1, za, zi, zb0, zm2])
    ck("SF9: no ratio between 1.90 and 2.05 on any convention (z or t)", not np.any((abs(allz) > 1.90) & (abs(allz) < 2.05)),
       allz)
    ck("SF10: no golden ratio within 0.3 of 1.96", min(abs(abs(zg) - 1.96)) >= 0.3, zg)
    ck("SF11: within the merged segment only 65-74 and 75 and over sit; 55-64 against 75 and over under one error",
       abs(med_gold[5] - med_gold[3]) < np.sqrt(res["groups"][3]["se"] ** 2 + res["groups"][5]["se"] ** 2), "")
    ck("SF12: ages top-coded at 85 in the person file", int(P["age"].max()) == 85, P["age"].max())
    ck("SF13: the weight is the same on a household's five rows", s["S"].groupby("hhid")["wgt"].nunique().max() == 1, "")
    singles = (P.groupby("hhid")["relationship"].apply(lambda r: not r.isin([2, 3]).any()))
    ck("SF14: households with no partner are placed by the reference person", np.array_equal(
        s["older"][singles.reindex(s["ids"]).to_numpy()], s["ref"][singles.reindex(s["ids"]).to_numpy()]), "")
    ck("SF15: the test is on the size of the difference", GOLD_SPLITS == segments(-zg), "")
    ck("SF16: the chart's intervals are the difference plus or minus 1.96 errors", all(
        abs(p["ci"][1] - p["diff"] - 1.96 * p["se"]) < 1 for p in res["pairs"]), "")

    # ------------------------------------------------------------------ bins
    for g in range(6):
        for nm, v in (("median", res["groups"][g]["median"]), ("se", res["groups"][g]["se"])):
            ck("bin: %s %s at least $100 from a rounding edge" % (nm, K.BAND_SHORT[g]), abs(abs(v) % 1000 - 500) >= 100, v)
    for g in range(5):
        p = res["pairs"][g]
        ck("bin: difference %d at least $100 from a rounding edge" % g, abs(abs(p["diff"]) % 1000 - 500) >= 100, p["diff"])
        ck("bin: difference %d's error at least $100 from a rounding edge" % g, abs(p["se"] % 1000 - 500) >= 100, p["se"])
        ck("bin: ratio %d at least 0.01 from a one-decimal edge" % g, abs(abs(p["ratio"]) * 10 % 1 - 0.5) >= 0.01, p["ratio"])
    for g in K.OLDER:
        ck("bin: Ask B income %s at least $150 from a rounding edge" % K.BAND_SHORT[g],
           abs(res["B"][g]["income"] % 1000 - 500) >= 150, res["B"][g]["income"])

    # ------------------------------------------------------------------ the asks: stops and the cracker sheet
    F = T.Files(target)
    gold_a = T.ask_a(F, **T.A_STOPS["golden"])
    gold_b = T.ask_b(F, **T.B_STOPS["golden"])
    for g in K.OLDER:
        ck("ask route: Ask A golden route %s equals the golden" % K.BAND_SHORT[g],
           (gold_a[g]["households"], G.thousands(gold_a[g]["balance"]), G.thousands(gold_a[g]["payments"])) ==
           (res["A"][g]["households"], G.thousands(res["A"][g]["balance"]), G.thousands(res["A"][g]["payments"])), "")
        ck("ask route: Ask B golden route %s equals the golden" % K.BAND_SHORT[g],
           (gold_b[g]["households"], G.thousands(gold_b[g]["income"])) ==
           (res["B"][g]["households"], G.thousands(res["B"][g]["income"])), "")
    st["stops"] = {}
    for nm, kw in T.A_STOPS.items():
        if nm == "golden":
            continue
        v = T.ask_a(F, **kw)
        st["stops"]["A " + nm] = v
        moved = [g for g in K.OLDER if (v[g]["households"], G.thousands(v[g]["balance"]), G.thousands(v[g]["payments"]))
                 != (res["A"][g]["households"], G.thousands(res["A"][g]["balance"]), G.thousands(res["A"][g]["payments"]))]
        ck("ask stop: Ask A '%s' lands off the golden in every group" % nm, moved == K.OLDER, (nm, moved))
    cr = T.ask_a(F, **T.A_STOPS["cracker"])
    for g in K.OLDER:
        for key in ("households", "balance", "payments"):
            a, b = cr[g][key], res["A"][g][key]
            same = (a == b) if key == "households" else (G.thousands(a) == G.thousands(b))
            ck("cracker sheet: Ask A %s %s missed" % (key, K.BAND_SHORT[g]), not same, (a, b))
    for nm, kw in T.B_STOPS.items():
        if nm == "golden":
            continue
        v = T.ask_b(F, **kw)
        st["stops"]["B " + nm] = v
        moved = [g for g in K.OLDER if (v[g]["households"], G.thousands(v[g]["income"])) !=
                 (res["B"][g]["households"], G.thousands(res["B"][g]["income"]))]
        ck("ask stop: Ask B '%s' lands off the golden in every group" % nm, moved == K.OLDER, (nm, moved))
    cb = T.ask_b(F, **T.B_STOPS["cracker"])
    for g in K.OLDER:
        ck("cracker sheet: Ask B households %s missed" % K.BAND_SHORT[g], cb[g]["households"] != res["B"][g]["households"], "")
        ck("cracker sheet: Ask B income %s missed" % K.BAND_SHORT[g], G.thousands(cb[g]["income"]) !=
           G.thousands(res["B"][g]["income"]), "")
    # each device has rows in every group it is ledgered to move
    pend = B.pending
    grp_rnc = {_group_of_survivor(B, k) for k in pend["rnc"]}
    ck("device H1: December deaths with a non-client reversionary reach at least two groups", len(grp_rnc & set(K.OLDER)) >= 2,
       grp_rnc)
    ck("device H1: December deaths without a reversionary in the extract", len(pend["db"]) >= 5, len(pend["db"]))
    ck("device H2: January deaths among pension members and partners", len(B.jan_deaths) >= 8, len(B.jan_deaths))
    ck("device H4: deferred annuities not yet paying", sum(1 for c in B.ann if c["pay_from"] and c["pay_from"] > K.AS_AT)
       >= 30, "")
    ck("device H3: January anniversaries increased before the pull", sum(1 for c in B.ann if I.jan_indexed(c) and
                                                                            c["index"] != "Level") >= 15, "")
    ck("device D6: reversionary contracts with both lives alive", sum(1 for c in B.ann if c["life2"] is not None and
                                                                       I.alive(B, c["life1"]) and I.alive(B, c["life2"])) >= 100, "")
    # the ask's conventions closed (judge r2): the partner a nomination names, a nominee who has died, the pull date
    dct = texts[K.DICT]
    ck("ask fork: the dictionary defines a Spouse or De facto partner nominee as the member's partner",
       "Spouse (the member's husband or wife), De facto partner (a person the member lives with as a couple)" in dct, "")
    ck("ask fork: the dictionary says the fund does not update nominee details",
       "the fund does not update them" in dct, "")
    ck("ask fork: the dictionary counts a household's partners alive on the date, clients or not",
       "partners at any date, our clients or not, are those alive on that date" in dct, "")
    ck("ask fork: the dictionary says where a partner who is not a client is known",
       "who is not our client is known to us only from these nominations, whatever their type" in dct, "")
    ck("ask fork: no annuity anniversary falls on the pull date", not any(
        (c["commenced"].month, c["commenced"].day) == (K.PULL.month, K.PULL.day) for c in B.ann), "")
    sp = F.sp.merge(F.acc[["account_no", "member_id"]], on="account_no")
    sp["owner"] = sp["member_id"].map(F.resolve)
    single = sp[sp["owner"].map(lambda x: len(F.members[F.by.at[x, "household_id"]]) == 1)]
    ck("ask fork: one partner per single-client household across its nominations",
       len(single) > 100 and (single.groupby("owner")["nominee_date_of_birth"].nunique() <= 1).all(), len(single))
    ck("ask fork: a partner nominee in a single-client household only for a married or de facto client",
       single["owner"].map(lambda x: F.by.at[x, "marital_status"]).isin(["Married", "De facto"]).all(), "")

    # ------------------------------------------------------------------ separation and distractors
    tmp = tempfile.mkdtemp(dir=os.path.dirname(target), suffix=".sep")
    try:
        for f in files:
            if f not in meta["ask_files"]:
                shutil.copy(os.path.join(target, f), tmp)
        s2 = G.load_survey(tmp)
        g2, p2, segs2 = G.main_call(E.Survey(s2["nw"], s2["w"], s2["R"]), K.band_of(s2["older"]))
        ck("separation: the call computed with every insurer file removed is the same", segs2 == res["segs"], segs2)
        for f in meta["distractor_files"]:
            os.remove(os.path.join(tmp, f)) if os.path.exists(os.path.join(tmp, f)) else None
        for f in meta["ask_files"]:
            shutil.copy(os.path.join(target, f), tmp)
        r3 = G.compute(tmp)
        ck("distractors: deleting both leaves every figure", r3["figures"] == fig, "")
    finally:
        shutil.rmtree(tmp)
    ck("separation: no survey column in an insurer file", not any(
        "hhid" in open(os.path.join(target, f), encoding="utf-8").readline() for f in (K.CRM, K.RIA, K.NOMS, K.ANN)), "")
    ck("separation: the brief carries no ask device word", not re.search(r"nominat|revers|deferred|anniversar|merged|deceased|died",
                                                                       texts[K.BRIEF], re.I), "")
    bul = texts[K.BULLETIN]
    ck("distractor: the bulletin is by the reference person's age", "age of the household reference person" in bul, "")
    ck("distractor: no bulletin age band shares a boundary with the brief's groups, so no structure for them reads off it",
       not re.search(r"\b(?:Under 35|35 to 44|45 to 54|55 to 64|65 to 74|75 and over|65 and over|55 and over)\b", bul), "")

    # ------------------------------------------------------------------ the goldens against figures.json
    gx = xlsx_text(os.path.join(golden, K.DELIVERABLES[1]))
    gd = docx_text(os.path.join(golden, K.DELIVERABLES[0]))
    for g in K.OLDER:
        for key, val in (("households", res["A"][g]["households"]), ("balance", G.thousands(res["A"][g]["balance"]) * 1000),
                         ("payments", G.thousands(res["A"][g]["payments"]) * 1000)):
            ck("golden xlsx: Ask A %s %s" % (key, K.BAND_SHORT[g]), str(val) in gx, val)
        ck("golden docx: Ask A balance %s" % K.BAND_SHORT[g], fig["A_balance_%d" % g] in gd, fig["A_balance_%d" % g])
        ck("golden docx: Ask B income %s" % K.BAND_SHORT[g], fig["B_income_%d" % g] in gd, fig["B_income_%d" % g])
    for g in range(5):
        ck("golden docx: difference %d and its error" % g, fig["diff_%d" % g] in gd and fig["se_diff_%d" % g] in gd, g)
    for g in range(6):
        ck("golden docx: median %d" % g, fig["median_%d" % g] in gd, fig["median_%d" % g])
    ck("golden docx: the five segments named", all(n in gd for n in res["names"]), res["names"])
    svg = open(os.path.join(golden, K.DELIVERABLES[2]), encoding="utf-8").read()
    ck("golden svg: no plotting library named", "matplotlib" not in svg.lower(), "")
    ck("golden svg: five pairs labelled", sum(1 for p in res["pairs"] if "%s / %s" % (p["lo"], p["hi"]) in svg) == 5, "")
    ck("figures.json: the call", fig["call"].startswith("Five segments"), fig["call"])

    # ------------------------------------------------------------------ realism sweep
    for f in files:
        b = open(os.path.join(target, f), "rb").read()
        ck("H1: %s carries no writing library's name" % f, not re.search(rb"reportlab|openpyxl|python-docx|matplotlib|"
                                                                          rb"ReportLab", b), f)
        if f.endswith((".xlsx", ".docx")):
            core = zipfile.ZipFile(os.path.join(target, f)).read("docProps/core.xml").decode()
            years = re.findall(r"<dcterms:(?:created|modified)[^>]*>(\d{4})", core)
            ck("H1: %s saved in 2025 to 2027" % f, years and all(2025 <= int(y) <= 2027 for y in years), years)
    for f, t in texts.items():
        ck("realism: no em or en dash in %s" % f, "—" not in t and "–" not in t, f)
    for f in files:
        if f.endswith(".csv"):
            t = open(os.path.join(target, f), encoding="utf-8").read()
            ck("realism: no em or en dash in %s" % f, "—" not in t and "–" not in t, f)
    crm = F.crm
    ck("realism: no client register date after the pull", pd.to_datetime(crm["client_since"]).max() <= pd.Timestamp(K.PULL)
       and crm["date_of_death"].max() <= pd.Timestamp(K.PULL), "")
    ck("realism: no annuity starts after the pull", F.ann["commenced"].max() <= pd.Timestamp(K.PULL), "")
    ck("realism: payments start after the pull only for deferred annuities", (F.ann.loc[F.ann.payments_start >
                                                                                       pd.Timestamp(K.PULL), "payments_start"]
                                                                               > F.ann.loc[F.ann.payments_start >
                                                                                           pd.Timestamp(K.PULL), "commenced"]).all(), "")
    ck("realism: pension balances not round", (F.acc["balance_31_12_2026"] % 1000 == 0).mean() < 0.01, "")
    ck("realism: group counts all differ", len({res["A"][g]["households"] for g in K.OLDER}) == 3, "")
    persons = re.findall(r"\b(Tracey Zimmerman|Ryan Ferrell|Wendy Koch)\b", " ".join(texts.values()))
    ck("realism: the card's personas appear", set(persons) >= {"Tracey Zimmerman", "Ryan Ferrell", "Wendy Koch"}, persons)
    card = json.load(open(os.path.join(REPO, ".claude/skills/fingerprint/cards/task261.json")))
    ck("realism: personas are the card's", set(card["world"]["people"]) == {"Tracey Zimmerman", "Ryan Ferrell", "Wendy Koch"},
       card["world"]["people"])
    ck("H21: invented names clear of real organisations", _h21(), "")
    print("checks: %d assertions passed" % ck.n)
    return ck.n


def _older_area(d):
    za = K.rng("older-areas").normal(0, 1, (K.STRATA, K.PSU_PER_STRATUM, 2))
    b = K.band_of(d["older"])
    sd = K.OLDER_AREA
    return np.where(b == 4, sd[0] * za[d["h"], d["p"], 0], np.where(b == 5, sd[1] * za[d["h"], d["p"], 1], 0.0))


def _group_of_survivor(B, k):
    nom = I.reversionary_of(B, k)
    return int(K.band_of(I.age_at(B.P[nom["person"]]["dob"])))


def _text(path):
    if path.endswith(".pdf"):
        return pdf_text(path)
    if path.endswith(".docx"):
        return docx_text(path)
    if path.endswith(".xlsx"):
        return xlsx_text(path)
    return open(path, encoding="utf-8").read()


REAL = ["TAL", "AIA", "MLC", "Zurich", "MetLife", "Resolution Life", "Challenger", "Generation Life", "ClearView",
        "Integrity Life", "NobleOak", "Hannover", "Swiss Re", "AMP", "Colonial First State", "Allianz", "Lifetime Financial",
        "Australian Retirement Trust", "AustralianSuper", "UniSuper", "Aware Super", "Hostplus", "Cbus", "Rest", "HESTA",
        "Insignia", "IOOF", "Macquarie", "BT", "Mercer", "Netwealth", "HUB24", "Kelvinator", "Kelvedon",
        "Melbourne Institute", "HILDA", "Household Income and Labour Dynamics", "Survey of Income and Housing",
        "Australian Bureau of Statistics", "Grattan Institute", "Centre for Independent Studies",
        "Household Finance and Consumption Survey", "Survey of Consumer Finances"]


def _h21():
    import difflib
    invented = ["Kelvane Life", "Kelvane Super", "Household Finance Research Centre", "National Household Wealth Survey"]
    for inv in invented:
        a = inv.lower().replace(" life", "").replace(" super", "")
        for r in REAL:
            b = r.lower()
            if a == b or (len(b) > 4 and (b in a or a in b)):
                return False
            if difflib.SequenceMatcher(None, a, b).ratio() > 0.85:
                return False
    return True
