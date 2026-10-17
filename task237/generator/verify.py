"""task237 verifier: rebuilds every owed figure from the pack with its own code (plain csv reading, its own valid-complete
rules, its own factor model fit) and compares it with the golden CSV, figures.json and submission.md."""
import csv
import datetime as dt
import json
import os
import re
import sys

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import minimize

CODES = ["BCS%02d" % i for i in range(1, 13)]
NEG = [1, 3, 5, 7, 9, 11]
POS = [0, 2, 4, 6, 8, 10]


def rows(path):
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def register(path):
    ws = load_workbook(path, read_only=True)["Register"]
    it = ws.iter_rows(values_only=True)
    head = next(it)
    return {r[0]: dict(zip(head, r)) for r in it if r[0]}


def score(answers):
    """TAS-4 s.5 on one respondent's raw answers (strings), None when three or more are unanswered"""
    s = []
    for i, a in enumerate(answers):
        a = int(a)
        if a in (9, -8, -9):
            continue
        s.append(6 - a if i in NEG else a)
    if len(s) < 10:
        return None, None
    item = [None if int(a) in (9, -8, -9) else (6 - int(a) if i in NEG else int(a)) for i, a in enumerate(answers)]
    return sum(s) / len(s) * 12, item


def online(target, files, reg):
    led = rows(os.path.join(target, files["ledger"]))
    member = {r["respondent_id"]: r["panel_member_id"] for r in led if r["respondent_id"]}
    merged = {r["panel_member_id"]: r["merged_into"] for r in led if r["outcome"] == "merged"}

    def root(a):
        while a in merged:
            a = merged[a]
        return a
    best = {}
    for r in rows(os.path.join(target, files["online"])):
        code = r["survey_code"].split("CCS6-BCS-")[1]
        v = {k.upper(): k for k in reg}[code]
        if r["status"] not in ("complete", "partial") or int(r["age"]) < 18 or any(r[c] == "" for c in CODES):
            continue
        if r["ui_language_at_end"] != v:
            continue
        start = dt.datetime.strptime(r["start_utc"], "%Y-%m-%dT%H:%M:%SZ") - dt.timedelta(hours=5)
        if not (str(reg[v]["pilot_start"])[:10] <= start.date().isoformat() <= str(reg[v]["pilot_end"])[:10]):
            continue
        t = r["end_utc"] or r["start_utc"]
        key = (v, root(member[r["respondent_id"]]))
        if key not in best or t < best[key][0]:
            best[key] = (t, [r[c] for c in CODES])
    return [(v, a) for (v, _), (_, a) in best.items()]


def capi(target, files, reg):
    checked = {}
    for r in rows(os.path.join(target, files["bc"])):
        checked[(r["case_id"], r["interviewer_id"])] = r["outcome"]
    iv = {}
    for r in rows(os.path.join(target, files["capi"])):
        k = (r["case_id"], r["version"], r["interviewer_id"], r["interview_date"])
        iv.setdefault(k, {})[r["variable"]] = r["value"]
    out = []
    for (case, v, who, day), d in iv.items():
        if any(c not in d for c in CODES) or int(d["AGE"]) < 18:
            continue
        if not (str(reg[v]["pilot_start"])[:10] <= day <= str(reg[v]["pilot_end"])[:10]):
            continue
        if checked.get((case, who)) == "not confirmed":
            continue
        out.append((v, [d[c] for c in CODES]))
    return out


def bifactor_share(S):
    """maximum likelihood bifactor fit written afresh: loadings on a general factor (all items) and a wording factor (the
    negatively worded items), log unique variances; the general factor's share of the total score's model variance"""
    p = 12

    def unpack(x):
        g = x[:12]
        w = np.zeros(12); w[NEG] = x[12:18]
        psi = np.exp(x[18:30])
        return g, w, psi

    def f(x):
        g, w, psi = unpack(x)
        Sig = np.outer(g, g) + np.outer(w, w) + np.diag(psi)
        try:
            L = np.linalg.cholesky(Sig)
        except np.linalg.LinAlgError:
            return 1e10
        logdet = 2 * np.log(np.diag(L)).sum()
        return logdet + np.trace(np.linalg.solve(Sig, S))
    d = np.sqrt(np.diag(S))
    x0 = np.concatenate([0.6 * d, 0.3 * d[NEG], np.log(0.4 * np.diag(S))])
    r = minimize(f, x0, method="BFGS", options={"maxiter": 5000, "gtol": 1e-8})
    g, w, psi = unpack(r.x)
    G, Wd = abs(g.sum()), abs(w.sum())
    return G ** 2 / (G ** 2 + Wd ** 2 + psi.sum())


def alpha(S):
    return 12 / 11 * (1 - np.trace(S) / S.sum())


def band(x):
    return "release" if x >= 0.80 else "hold" if x >= 0.70 else "rework"


def main(target, golden, figures_path, submission_path):
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import build as B
    import golden as GD
    reg = register(os.path.join(target, B.FILES["reg"]))
    comp = online(target, B.FILES, reg) + capi(target, B.FILES, reg)
    order = list(reg)
    counts = {v: sum(1 for x, _ in comp if x == v) for v in order}
    means, calls, alphas = {}, {}, {}
    for v in order:
        sc = [score(a) for x, a in comp if x == v]
        tot = [s for s, _ in sc if s is not None]
        means[v] = sum(tot) / len(tot)
        full = np.array([it for s, it in sc if it is not None and None not in it], dtype=float)
        S = np.cov(full, rowvar=False)
        calls[v] = band(bifactor_share(S))
        alphas[v] = alpha(S)
    gold = {r["version"]: r for r in rows(os.path.join(golden, GD.CSV))}
    assert list(gold) == order, "golden rows in the register's order"
    for v in order:
        g = gold[v]
        assert int(g["valid_pilot_completes"]) == counts[v], (v, g["valid_pilot_completes"], counts[v])
        assert g["pilot_mean_score"] == "%.2f" % (np.floor(means[v] * 100 + 0.5) / 100), (v, g["pilot_mean_score"], means[v])
        assert g["decision"] == calls[v], (v, g["decision"], calls[v])
    fig = json.load(open(figures_path))
    rows_f = {r["version"]: r for r in fig["deliverables"][GD.CSV]}
    assert all(rows_f[v]["valid_pilot_completes"] == counts[v] and rows_f[v]["decision"] == calls[v] for v in order)
    lang = {v: reg[v]["language"] for v in order}
    assert {lang[v] for v in order if calls[v] == "release"} == set(fig["call"]["release"])
    up = [alphas[v] for v in order if band(alphas[v]) == "release" and calls[v] != "release"]
    assert fig["rivals"]["alpha_range_released_wrongly"] == ["%.2f" % min(up), "%.2f" % max(up)]
    sub = open(submission_path).read()
    for v in order:
        line = "%s (%s): %s, %s valid pilot completes, pilot mean score %s" % (
            lang[v], v, calls[v], format(counts[v], ","), gold[v]["pilot_mean_score"])
        assert line in sub, line
    assert "at %.2f to %.2f would release" % (min(up), max(up)) in sub
    print("verify: %d versions, every count, mean and decision rebuilt independently" % len(order))
