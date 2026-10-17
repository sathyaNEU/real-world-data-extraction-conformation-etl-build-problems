"""The golden deliverables: respiratory_feed_escalations_2027-W03.xlsx and respiratory_feed_completeness_2027-W03.png."""
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import params as P
import postprocess

XLSX = "respiratory_feed_escalations_2027-W03.xlsx"
PNG = "respiratory_feed_completeness_2027-W03.png"
NAVY = "1F3A5F"
WARN = "B03A2E"
HEAD = PatternFill("solid", fgColor=NAVY)
ESC = PatternFill("solid", fgColor="F6DCD8")
LINE = Border(bottom=Side(style="thin", color="C9CED6"))
AUTHOR = P.PEOPLE["manager"]


def _head(ws, row, labels, widths):
    for j, (lab, w) in enumerate(zip(labels, widths), 1):
        c = ws.cell(row=row, column=j, value=lab)
        c.font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left" if j <= 3 else "right")
        ws.column_dimensions[chr(64 + j)].width = w
    ws.row_dimensions[row].height = 30


def _cell(ws, r, c, v, fmt=None, bold=False, color=None, fill=None):
    x = ws.cell(row=r, column=c, value=v)
    x.font = Font(name="Calibri", size=10, bold=bold, color=color)
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = fill
    x.border = LINE
    x.alignment = Alignment(horizontal="left" if c <= 3 else "right")
    return x


def workbook(path, A):
    wb = Workbook()
    ws = wb.active
    ws.title = "Escalations"
    L = A["labname"]
    ws["A1"] = "Escalate %s and %s, all three streams; no other stream" % (L["L03"], L["L10"])
    ws["A1"].font = Font(name="Calibri", bold=True, size=13, color=NAVY)
    ws["A2"] = ("Respiratory feed escalations under FCP-2 for MMWR week 3 of 2027 (17 to 23 January), set Monday 25 January "
                "2027. Completeness is the week's tests with results on the platform by 08:00 on 25 January over the tests "
                "each stream's week-3 positivity predicts its laboratory ran; the week-3 aggregate reports are due "
                "Wednesday 27 January.")
    ws["A2"].font = Font(name="Calibri", italic=True, size=9, color="444444")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:H2")
    ws.row_dimensions[2].height = 40
    _head(ws, 4, ["Laboratory", "Lab", "Stream", "Tests on the platform", "Positivity (%)", "Tests predicted from positivity",
                  "Completeness (%)", "Escalate"], [44, 7, 12, 13, 11, 15, 13, 10])
    r = 5
    for s in A["streams"]:
        f = ESC if s["escalate"] else None
        _cell(ws, r, 1, s["lab"], fill=f, bold=s["escalate"])
        _cell(ws, r, 2, s["lab_id"], fill=f)
        _cell(ws, r, 3, s["virus"], fill=f)
        _cell(ws, r, 4, s["received"], "#,##0", fill=f)
        _cell(ws, r, 5, s["positivity"], "0.0", fill=f)
        _cell(ws, r, 6, s["est"], "#,##0", fill=f)
        _cell(ws, r, 7, s["comp"], "0", fill=f, bold=s["escalate"], color=WARN if s["escalate"] else None)
        _cell(ws, r, 8, "Yes" if s["escalate"] else "No", fill=f, bold=s["escalate"], color=WARN if s["escalate"] else None)
        r += 1
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = "A4:H%d" % (r - 1)
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0

    w2 = wb.create_sheet("Week 2")
    w2["A1"] = "Week 2 completeness (10 to 16 January 2027), now that the week-2 aggregate reports are in"
    w2["A1"].font = Font(name="Calibri", bold=True, size=12, color=NAVY)
    w2["A2"] = ("Tests of week 2 whose first result reached the platform by 08:00 county time on Monday 18 January, over the "
                "tests in the laboratory's week-2 aggregate report of record (FCP-2 sections 3 and 5).")
    w2["A2"].font = Font(name="Calibri", italic=True, size=9, color="444444")
    _head(w2, 4, ["Laboratory", "Lab", "Stream", "Tests in the aggregate report", "Results by the cutoff",
                  "Completeness (%)"], [44, 7, 12, 15, 13, 14])
    r = 5
    for w in A["week2"]:
        _cell(w2, r, 1, w["lab"]); _cell(w2, r, 2, w["lab_id"]); _cell(w2, r, 3, w["virus"])
        _cell(w2, r, 4, w["report"], "#,##0"); _cell(w2, r, 5, w["received"], "#,##0")
        _cell(w2, r, 6, w["comp"], "0.0")
        r += 1
    w2.freeze_panes = "D5"

    sp = wb.create_sheet("Specimens")
    sp["A1"] = "Specimens tested in week 3 with results on the platform"
    sp["A1"].font = Font(name="Calibri", bold=True, size=12, color=NAVY)
    sp["A2"] = ("One specimen is one collection from one patient (FCP-2 section 6): the laboratory's patient number and the "
                "collection time, whatever accession numbers its tests carry.")
    sp["A2"].font = Font(name="Calibri", italic=True, size=9, color="444444")
    _head(sp, 4, ["Laboratory", "Lab", "Specimens"], [44, 7, 12])
    r = 5
    for s in A["specimens"]:
        _cell(sp, r, 1, s["lab"]); _cell(sp, r, 2, s["lab_id"]); _cell(sp, r, 3, s["specimens"], "#,##0")
        sp.cell(row=r, column=3).alignment = Alignment(horizontal="right")
        r += 1
    _cell(sp, r, 1, "All ten laboratories", bold=True); _cell(sp, r, 2, "")
    _cell(sp, r, 3, sum(s["specimens"] for s in A["specimens"]), "#,##0", bold=True)
    sp.cell(row=r, column=3).alignment = Alignment(horizontal="right")

    nt = wb.create_sheet("Notes")
    notes = [
        "Notes",
        "As of 08:00 county time (America/Chicago) on Monday 25 January 2027: elr_respiratory_results_extract_2027-01-25.csv "
        "and lab_weekly_aggregate_reports.csv (reports of record, 2022 week 40 to 2027 week 2).",
        "Tests follow FCP-2 section 2: one per laboratory, accession and virus; RV-FLUA and RV-FLUB on one accession are one "
        "influenza test; a corrected result stands in place of the result it corrects; cancelled orders are not tests; the "
        "week is the county date of the first release.",
        "Tests predicted from positivity: for each stream, the log of the weekly tests regressed on the week's positivity and a "
        "linear time trend over every report of record, leaving out the weeks holding Thanksgiving, Christmas Day and New "
        "Year's Day, read at the stream's week-3 positivity. A lost result message takes a positive and a negative alike, so "
        "the positivity of the results that arrived is the laboratory's.",
        "Completeness above 100 means the stream's results ran ahead of the prediction; the escalation line is 80 per cent, "
        "compared unrounded (FCP-2 section 4).",
        "Receipt times: the platform wrote them in county time until the state hub's broker migration at 06:00 on 19 January "
        "2027 (CHG-2027-0019) and in UTC from it; week 2's cutoff is read on that basis.",
        "Prepared by %s, surveillance data manager, for %s, laboratory liaison." % (AUTHOR, P.PEOPLE["liaison"]),
    ]
    for i, t in enumerate(notes, 1):
        c = nt.cell(row=i, column=1, value=t)
        c.font = Font(name="Calibri", size=10, bold=(i == 1), color=NAVY if i == 1 else None)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    nt.column_dimensions["A"].width = 110
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.save(path)
    postprocess.scrub(path, AUTHOR, "2027-01-25")
    postprocess.finish_office(path, "2027-01-25T14:05:00Z", "2027-01-25T16:40:00Z", P.DEPT)


def _short(name):
    return name.replace(" Laboratory Services", "").replace(" Laboratory", "")


def chart(path, A):
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 8.5})
    fig, ax = plt.subplots(figsize=(9.0, 10.2), dpi=150)
    ss = A["streams"][::-1]
    y, labels, gap, ypos = [], [], 0.0, 0.0
    prev = None
    for s in ss:
        if prev is not None and s["lab_id"] != prev:
            ypos += 0.6
        y.append(ypos); ypos += 1
        labels.append("%s  %s" % (_short(s["lab"]), s["virus"]))
        prev = s["lab_id"]
    vals = [100 * s["comp_raw"] for s in ss]
    colors = ["#" + WARN if s["escalate"] else "#8EA3B9" for s in ss]
    ax.barh(y, vals, color=colors, height=0.72, zorder=3)
    ax.set_yticks(y)
    ax.set_yticklabels(labels)
    ax.axvline(80, color="#" + WARN, linestyle="--", linewidth=1.2, zorder=4)
    ax.text(80.8, y[-1] + 0.9, "escalation line 80%", color="#" + WARN, fontsize=8, va="bottom")
    for yy, s, v in zip(y, ss, vals):
        if s["escalate"]:
            ax.text(v + 1.0, yy, "%d%%" % s["comp"], va="center", ha="left", fontsize=8, color="#" + WARN,
                    fontweight="bold")
    ax.set_xlim(0, 112)
    ax.set_ylim(-0.8, y[-1] + 1.8)
    ax.set_xticks(range(0, 111, 10))
    ax.set_xticklabels(["%d%%" % t for t in range(0, 111, 10)])
    ax.grid(axis="x", color="#E5E5E5", zorder=0)
    ax.set_xlabel("Week-3 completeness: tests on the platform over the tests the stream's positivity predicts")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    fig.tight_layout(rect=(0, 0, 1, 0.94))
    L = A["labname"]
    fig.text(0.015, 0.988, "Wexmoor and Westmarch are short of results in all three streams; escalate both",
             fontsize=11.5, fontweight="bold", va="top")
    fig.text(0.015, 0.966, "Respiratory result streams, MMWR week 3 of 2027 (17 to 23 January), set Monday 25 January 2027 "
                           "before the week's aggregate reports.\nEvery other stream sits above the line at %d%% or more; "
                           "escalate below 80%% (FCP-2 section 4)." % A["min_rest"],
             fontsize=8.5, va="top", color="#333333", linespacing=1.4)
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)


def write(golden_dir, A):
    os.makedirs(golden_dir, exist_ok=True)
    workbook(os.path.join(golden_dir, XLSX), A)
    chart(os.path.join(golden_dir, PNG), A)
