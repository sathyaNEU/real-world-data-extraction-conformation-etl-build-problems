"""Independent verifier: reads only the shipped bundle (target/, golden/, figures.json, submission.md) and recomputes
every graded figure with its own code, then checks the goldens and the write-up carry them."""
import datetime as dt
import json
import os
import re

import numpy as np
import pandas as pd
import statsmodels.api as sm
from openpyxl import load_workbook

R = "elr_respiratory_results_extract_2027-01-25.csv"
AGG = "lab_weekly_aggregate_reports.csv"
REG = "respiratory_stream_register.xlsx"
CODES = "elr_result_codes.csv"


def holiday(y, w):
    j4 = dt.date(y, 1, 4)
    start = j4 - dt.timedelta(days=(j4.weekday() + 1) % 7) + dt.timedelta(weeks=w - 1)
    for i in range(7):
        d = start + dt.timedelta(days=i)
        if (d.month, d.day) in ((12, 25), (1, 1)) or (d.month == 11 and d.weekday() == 3 and 22 <= d.day <= 28):
            return True
    return False


def tests_of(target):
    res = pd.read_csv(os.path.join(target, R), dtype=str, keep_default_na=False)
    res = res.rename(columns={"released_utc": "released", "specimen_collected_utc": "specimen_collected"})
    labs = pd.read_excel(os.path.join(target, REG), sheet_name="Laboratories", dtype=str)
    codes = pd.read_csv(os.path.join(target, CODES), dtype=str)
    fac = {f.strip(): r.lab_id for r in labs.itertuples() for f in r.sending_facility_ids.split(",")}
    vir = dict(zip(codes.test_code, codes.virus))
    d = res[res.result_status.isin(["F", "C"])].copy()
    d["lab"] = d.sending_facility.map(fac)
    d["virus"] = d.test_code.map(vir)
    rel = pd.to_datetime(d.released) - pd.Timedelta(hours=6)
    rec = pd.to_datetime(d.received)
    switch = pd.Timestamp("2027-01-19 06:00")
    d["rec"] = np.where(rec < switch, rec, rec - pd.Timedelta(hours=6))
    d["rel"] = rel
    d["order"] = np.where(d.result_status == "C", 1, 0)
    d = d.sort_values(["lab", "accession", "test_code", "rel", "order"])
    stand = d.drop_duplicates(["lab", "accession", "test_code"], keep="last")
    pos = stand.assign(p=stand.result.eq("DETECTED")).groupby(["lab", "virus", "accession"])["p"].max()
    g = d.groupby(["lab", "virus", "accession"])
    t = pd.DataFrame({"rel": g["rel"].min(), "rec": g["rec"].min(), "pat": g["patient_mrn"].first(),
                      "coll": g["specimen_collected"].first(), "pos": pos}).reset_index()
    t["week"] = ((t.rel.dt.normalize() - pd.Timestamp("2027-01-03")).dt.days // 7) + 1
    return t, labs


def poisson(Xm, y):
    return sm.GLM(y, Xm, family=sm.families.Poisson()).fit(tol=1e-12, maxiter=200).params


def main(target, golden, figs_path, sub_path):
    t, labs = tests_of(target)
    agg = pd.read_csv(os.path.join(target, AGG))
    rep = agg.sort_values("version").drop_duplicates(["lab_id", "virus", "mmwr_year", "mmwr_week"], keep="last")
    rep = rep.sort_values(["mmwr_year", "mmwr_week"])
    weeks = sorted({(int(a), int(b)) for a, b in zip(rep.mmwr_year, rep.mmwr_week)})
    idx = {w: i for i, w in enumerate(weeks)}
    n = len(weeks)
    tt = np.arange(n + 1) / 52.18
    keep = np.array([not holiday(*w) for w in weeks])
    w3 = t[t.week == 3]
    w2 = t[t.week == 2]
    cut = pd.Timestamp("2027-01-18 08:00")
    assert weeks[-2:] == [(2027, 1), (2027, 2)]
    assert pd.to_datetime(rep[(rep.mmwr_year < 2027) | (rep.mmwr_week < 2)].submitted).max() < cut
    got, miss = {}, {"golden": [], "received": [], "persist": []}
    for (lab, v), g in rep.groupby(["lab_id", "virus"]):
        T = np.zeros(n); Pp = np.zeros(n)
        for r in g.itertuples():
            T[idx[(r.mmwr_year, r.mmwr_week)]] = r.tests; Pp[idx[(r.mmwr_year, r.mmwr_week)]] = r.positives
        assert T.min() > 0
        x = w3[(w3.lab == lab) & (w3.virus == v)]
        rec, pos = len(x), int(x.pos.sum())
        b = poisson(np.c_[np.ones(keep.sum()), Pp[keep] / T[keep], tt[:n][keep]], T[keep])
        est = float(np.exp(b[0] + b[1] * pos / rec + b[2] * tt[n]))
        got[(lab, v)] = dict(rec=rec, pos=pos, est=est, comp=rec / est)
        # FCP-2 section 5: the same way made at week 2's cutoff, from the reports before week 2
        h = w2[(w2.lab == lab) & (w2.virus == v) & (w2.rec <= cut)]
        k1 = keep[:n - 1]
        b2 = poisson(np.c_[np.ones(k1.sum()), Pp[:n - 1][k1] / T[:n - 1][k1], tt[:n - 1][k1]], T[:n - 1][k1])
        e2 = float(np.exp(b2[0] + b2[1] * h.pos.sum() / len(h) + b2[2] * tt[n - 1]))
        miss["golden"].append(abs(e2 / T[n - 1] - 1))
        miss["received"].append(abs(len(h) / T[n - 1] - 1))
        miss["persist"].append(abs(T[n - 2] / T[n - 1] - 1))
    check = {k: round(float(100 * max(m)), 1) for k, m in miss.items()}
    assert check["golden"] * 5 < min(check["received"], check["persist"]), check
    esc = {k for k, v in got.items() if v["comp"] < 0.80}
    assert esc == {("L03", v) for v in ("Influenza", "RSV", "SARS-CoV-2")} | {("L10", v) for v in ("Influenza", "RSV", "SARS-CoV-2")}, esc
    # week 2
    inby = w2[w2.rec <= cut].groupby(["lab", "virus"]).size()
    r2 = rep[(rep.mmwr_year == 2027) & (rep.mmwr_week == 2)].set_index(["lab_id", "virus"]).tests
    wk2 = {k: 100.0 * inby[k] / r2[k] for k in r2.index}
    # specimens
    spec = w3.groupby("lab").apply(lambda g: len(set(zip(g.pat, g.coll)))).to_dict()

    # the golden workbook
    wb = load_workbook(os.path.join(golden, "respiratory_feed_escalations_2027-W03.xlsx"))
    ws = wb["Escalations"]
    rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]
    assert len(rows) == 29
    for r in rows:
        k = (r[1], r[2]); g = got[k]
        assert r[3] == g["rec"] and abs(r[4] - round(100 * g["pos"] / g["rec"], 1)) < 1e-9, (k, r)
        assert r[5] == int(round(g["est"])) and r[6] == round(100 * g["comp"]), (k, r)
        assert (r[7] == "Yes") == (k in esc), (k, r)
    ws = wb["Week 2"]
    rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]
    assert len(rows) == 29
    for r in rows:
        k = (r[1], r[2])
        assert r[3] == r2[k] and r[4] == inby[k] and abs(r[5] - round(wk2[k], 1)) < 1e-9, (k, r)
    ws = wb["Specimens"]
    rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0] and r[1]]
    assert len(rows) == 10 and all(r[2] == spec[r[1]] for r in rows), rows
    svg = os.path.join(golden, "respiratory_feed_completeness_2027-W03.svg")
    assert open(svg, "rb").read(200).lstrip().startswith(b"<?xml") and os.path.getsize(svg) > 20_000
    assert sorted(os.listdir(golden)) == ["respiratory_feed_completeness_2027-W03.svg", "respiratory_feed_escalations_2027-W03.xlsx"]

    # figures.json and the write-up
    figs = json.load(open(figs_path))
    for k, x in check.items():
        assert figs["week2_check_largest_miss"][k] == x, (k, x, figs["week2_check_largest_miss"])
    names = dict(zip(labs.lab_id, labs.laboratory))
    assert sorted(figs["call"]) == sorted("%s %s" % (names[k[0]], k[1]) for k in esc)
    for k, v in figs["week2"].items():
        lab, vir = k.split(" ", 1)
        assert abs(v["completeness"] - round(wk2[(lab, vir)], 1)) < 1e-9
    assert {names[k]: v for k, v in spec.items()} == figs["specimens"]
    sub = open(sub_path).read()
    assert "\u2014" not in sub and "\u2013" not in sub
    steps = sub.split("## 3. Step-by-Step Solution")[1].split("## 4.")[0]
    step3 = next(l for l in steps.splitlines() if l.startswith("3. "))
    assert ("missed some stream by %.1f per cent at most; those held tests themselves missed by %.1f and last week's "
            "report by %.1f." % (check["golden"], check["received"], check["persist"])) in step3, step3
    block4 = sub.split("## 4. Deliverable Answers")[1]
    for lab, name in names.items():
        line = next(l for l in block4.splitlines() if l.strip().startswith("- %s:" % name) and "influenza" in l and "." in l)
        for vir in ("Influenza", "RSV", "SARS-CoV-2"):
            if (lab, vir) in wk2:
                low = {"Influenza": "influenza"}.get(vir, vir)
                assert re.search(r"%s %.1f(;|$)" % (re.escape(low), round(wk2[(lab, vir)], 1)), line), (line, vir)
        assert ("- %s: %s" % (name, format(spec[lab], ","))) in block4, name
    print("verify: escalations, the week-2 check (%s), %d week-2 figures, %d specimen counts, goldens and write-up "
          "recomputed from the bundle" % (check, len(wk2), len(spec)))
