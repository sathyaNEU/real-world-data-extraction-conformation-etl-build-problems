"""task234 documents: the feed completeness protocol FCP-2 (pdf), the stream register (xlsx) and the platform change
log (csv)."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess

OFFICE = P.DEPT


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=13, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=14, leading=18, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=14, spaceBefore=8,
                        spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10, textColor=colors.HexColor("#444444"))
    return dict(body=body, h1=h1, h2=h2, small=small)


def _footer(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(18 * mm, 10 * mm, text)
        canvas.drawRightString(LETTER[0] - 18 * mm, 10 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _table(rows, widths, size=8.0):
    t = Table(rows, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", size), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", size),
                           ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6ECF2")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


# ------------------------------------------------------------------ FCP-2
def protocol(path):
    S = _styles(); b, h1, h2, sm = S["body"], S["h1"], S["h2"], S["small"]
    st = [Paragraph("FCP-2 Feed completeness protocol: respiratory laboratory results", h1),
          Paragraph("%s, %s. Version 2, in force from 1 October 2026. Approved 21 September 2026 by %s, respiratory "
                    "epidemiologist, who keeps this protocol. Replaces FCP-1 (2023)." % (P.DEPT, P.UNIT, P.PEOPLE["epi"]), sm),
          Spacer(1, 4),
          Paragraph("1. Scope", h2),
          Paragraph("This protocol applies to the result streams listed in the respiratory stream register from 1 October "
                    "2026. A result stream is one laboratory's electronic results for one of influenza, RSV and SARS-CoV-2 "
                    "sent to the county's surveillance platform. A laboratory that sends under more than one facility "
                    "identifier is one laboratory; the register lists its identifiers.", b),
          Paragraph("2. Tests", h2),
          Paragraph("A test is one virus tested on one accession. A laboratory that reports influenza A and influenza B as "
                    "separate results on an accession (test codes RV-FLUA and RV-FLUB) has run one influenza test on it. A "
                    "corrected result (status C) replaces the result it corrects, and the test is counted once. A cancelled "
                    "order (status X) is not a test. An inconclusive result is a test. A test belongs to the MMWR week of "
                    "the county date (America/Chicago) on which the laboratory first released its result.", b),
          Paragraph("3. Completeness", h2),
          Paragraph("A stream's completeness for an MMWR week is the number of the stream's tests of that week whose result "
                    "reached the platform by 08:00 county time on the Monday after the week, divided by the number of tests "
                    "the laboratory's weekly aggregate report counts for that virus and week, as a percentage.", b),
          Paragraph("Each laboratory files a weekly aggregate report of the tests it ran and the positives it found for each "
                    "virus in the MMWR week, every test whatever the patient's residence, by the Wednesday after the week. "
                    "A laboratory that corrects a report resubmits it; the latest submission is the report of record.", b),
          Paragraph("4. Escalation", h2),
          Paragraph("Every Monday the surveillance data manager sets the streams the laboratory liaison escalates for the "
                    "week just closed: each stream whose completeness for that week is below 80 per cent. The escalations "
                    "are set that Monday, before the week's aggregate reports are due, so that a laboratory losing results "
                    "is called while the week's results can still be resent. Completeness is compared with the line "
                    "unrounded.", b),
          Paragraph("The liaison calls each escalated laboratory the same day, asks it to check its interface and to resend "
                    "the week's missing results, and records the call in the stream's log.", b),
          Paragraph("5. Closing a week", h2),
          Paragraph("When a week's aggregate reports of record are in, the data manager records each stream's completeness "
                    "for the week, and the liaison quotes it to the laboratory at the next call. Results that reach the "
                    "platform after the Monday cutoff stay on the platform for case surveillance and do not count toward "
                    "the week's completeness.", b),
          Paragraph("The data manager then sets the week's reports beside the estimates of each stream's tests for the week "
                    "that could have been made at that Monday's cutoff. The next Monday's escalations are judged against "
                    "the week's tests estimated in the way whose largest miss across the streams, as a share of the tests "
                    "the reports counted, was smallest.", b),
          Paragraph("6. Specimens", h2),
          Paragraph("A specimen is one collection from one patient. The liaison opens each call with the number of the "
                    "laboratory's specimens tested in the week whose results are on the platform.", b),
          Paragraph("7. Roles", h2),
          _table([["Role", "Holder", "In this protocol"],
                  ["Surveillance data manager", P.PEOPLE["manager"], "sets the week's escalations (s.4); records each "
                                                                     "week's completeness (s.5)"],
                  ["Laboratory liaison", P.PEOPLE["liaison"], "calls the escalated laboratories (s.4)"],
                  ["Respiratory epidemiologist", P.PEOPLE["epi"], "keeps this protocol and the stream register"]],
                 [48 * mm, 36 * mm, 92 * mm]),
          ]
    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title="FCP-2 Feed completeness protocol", author=P.PEOPLE["epi"],
                            creator=P.PEOPLE["epi"], invariant=1)
    foot = "FCP-2 v2 | %s, %s" % (P.DEPT, P.UNIT)
    doc.build(st, onFirstPage=_footer(foot), onLaterPages=_footer(foot))
    postprocess.scrub(path, P.PEOPLE["epi"], "2026-09-21")


# ------------------------------------------------------------------ the stream register
def register(path, labs, streams):
    wb = Workbook(); ws = wb.active; ws.title = "Laboratories"
    hdr = list(labs.columns)
    ws.append(hdr)
    for r in labs.itertuples(index=False):
        ws.append(list(r))
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="E6ECF2")
    for i, w in enumerate([8, 44, 20, 26, 28, 26], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    wd = wb.create_sheet("Streams")
    wd.append(list(streams.columns))
    for r in streams.itertuples(index=False):
        wd.append(list(r))
    for c in wd[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="E6ECF2")
    for i, w in enumerate([10, 8, 13, 20], 1):
        wd.column_dimensions[get_column_letter(i)].width = w
    wd.freeze_panes = "A2"
    wb.save(path)
    postprocess.scrub(path, P.PEOPLE["epi"], "2026-10-26")
    postprocess.finish_office(path, "2026-09-08T15:02:00Z", "2026-10-26T13:41:00Z", OFFICE)


# ------------------------------------------------------------------ the platform change log
def change_log():
    E, M, L, S, IT = P.PEOPLE["epi"], P.PEOPLE["manager"], P.PEOPLE["liaison"], "State ELR hub", "IT Infrastructure"
    rows = [
        ("CHG-2026-0398", "2026-09-02 19:00", "Surveillance platform", "Respiratory season reports configured from MMWR "
         "2026 week 40", E),
        ("CHG-2026-0412", "2026-09-08 09:00", "Surveillance platform", "Respiratory stream register published for FCP-2: "
         "ten laboratories, 29 result streams", E),
        ("CHG-2026-0419", "2026-09-15 22:00", "ELR inbound", "Interface engine patched to 11.2.4", IT),
        ("CHG-2026-0431", "2026-10-01 00:00", "Surveillance platform", "FCP-2 in force for the respiratory result streams", E),
        ("CHG-2026-0452", "2026-10-14 07:00", "Surveillance platform", "Influenza-associated pediatric death case "
         "definition updated for 2026-27", E),
        ("CHG-2026-0466", "2026-10-26 07:30", "ELR inbound", "Corrigan Valley Hospital Laboratory: facility 05D2290465 "
         "(north campus laboratory, merged into Corrigan Valley on 1 October) mapped to the laboratory", L),
        ("CHG-2026-0474", "2026-11-02 18:00", "Surveillance platform", "Reporting dashboard moved to the county's new "
         "reporting server", IT),
        ("CHG-2026-0481", "2026-11-09 21:00", "ELR inbound", "Test code RV-COV2 description aligned with the state code list "
         "(no code change)", L),
        ("CHG-2026-0498", "2026-11-16 21:00", "ELR inbound", "TLS certificate renewed on the county's inbound ELR endpoint", IT),
        ("CHG-2026-0505", "2026-11-23 06:00", "Surveillance platform", "Storage expanded on the ELR staging database", IT),
        ("CHG-2026-0517", "2026-12-03 18:00", "Surveillance platform", "Nightly case extract moved from 02:00 to 03:30", IT),
        ("CHG-2026-0529", "2026-12-10 09:00", "Surveillance platform", "Respiratory season user access review; two accounts "
         "of staff who left removed", M),
        ("CHG-2026-0536", "2026-12-15 20:00", S, "Firewall rules opened for the hub's new cloud broker ahead of the "
         "migration", S),
        ("CHG-2027-0003", "2027-01-04 08:00", "Surveillance platform", "MMWR 2027 week calendar loaded", M),
        ("CHG-2027-0007", "2027-01-06 20:00", "Surveillance platform", "Platform patch 7.4.2: case management screens", IT),
        ("CHG-2027-0019", "2027-01-19 06:00", S, "State hub's message broker migrated to the hub's new cloud broker; the "
         "county's inbound endpoint repointed to it", S),
        ("CHG-2027-0024", "2027-01-21 14:00", "Surveillance platform", "Respiratory dashboard: SARS-CoV-2 variant panel "
         "added", M),
    ]
    return pd.DataFrame(rows, columns=["change_id", "implemented_at", "system", "change", "requested_by"])

