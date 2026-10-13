"""The two deliverables: elr_sla_thresholds_2026-27.csv (each feed's SLA lateness for the NCP 4.2 feed register, with its
expected secondary cases a year at that setting and at the next longer one) and elr_sla_basis_2026-27.xlsx (the basis:
each condition's secondary transmission, each feed's expected secondary cases at every allowed setting with a chart, each
feed's reports and late reports in the year to 30 September 2026 and the urgent cases it was first to notify)."""
import csv
import os

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.marker import DataPoint, Marker
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import params as P
import postprocess
import world as W

CSV = "elr_sla_thresholds_2026-27.csv"
WORKBOOK = "elr_sla_basis_2026-27.xlsx"
THIN = Side(style="thin", color="B9C2BD")
SETTINGS = P.SETTINGS
SHORT = {"CRBH": "Base hospital", "LPG": "Larkspur", "HDN": "Halden", "KES": "Kestrel", "HPS": "Highfield",
         "SERL": "Reference laboratory"}
COND_ORDER = ["Shigellosis", "Shiga toxin-producing E. coli infection", "Cryptosporidiosis"]
LINE = {"CRBH": "1F4E79", "LPG": "C55A11", "HDN": "548235", "KES": "7030A0", "HPS": "BF9000", "SERL": "7F7F7F"}


def write_csv(A, path):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["feed", "sending_facility", "sla_lateness_hours", "expected_secondary_cases_a_year_at_setting",
                    "expected_secondary_cases_a_year_at_next_longer_setting"])
        for f in P.FEEDS:
            x = A["feeds"][f]
            w.writerow([W.FEED[f]["name"], W.FEED[f]["codes"][-1][0], x["setting"], "%.2f" % x["at"],
                        "%.2f" % x["next"] if x["next"] is not None else ""])


def _header(ws, row, cols, numeric_from=2):
    for j, h in enumerate(cols, start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(name="Calibri", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5E4E")
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="right" if j >= numeric_from else None)
        c.border = Border(bottom=THIN)
    ws.row_dimensions[row].height = 30


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _top(ws, title, sub):
    ws["A1"] = title
    ws["A1"].font = Font(name="Calibri", bold=True, size=12)
    ws["A2"] = sub
    ws["A2"].font = Font(name="Calibri", italic=True, size=9, color="555555")
    ws.freeze_panes = "B5"
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _title(text, size, bold=True):
    cp = CharacterProperties(sz=size, b=bold)
    return Title(tx=Text(rich=RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp),
                                                    r=[RegularTextRun(rPr=cp, t=text)])])), overlay=False)


def _note(ws, row, text, ncols, width):
    ws.merge_cells("A%d:%s%d" % (row, get_column_letter(ncols), row))
    c = ws.cell(row=row, column=1, value=text)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    c.font = Font(name="Calibri", size=9, color="333333")
    ws.row_dimensions[row].height = 13.5 * max(1, -(-len(text) // int(width * 1.1)))


SUB = ("Basis: the notifications that opened the unit's childcare investigations, authorised 1 October 2023 to "
       "30 September 2026; tolerance 1.00 expected secondary cases a year (CD-ST-07 section 2). Communicable Disease "
       "Surveillance, 16 November 2026.")


def workbook(A, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Settings"
    title = "SLA lateness for NCP 4.2: " + ", ".join("%s %d h" % (SHORT[f], A["feeds"][f]["setting"]) for f in P.FEEDS)
    _top(ws, title, SUB)
    cols = ["Feed", "Sending facility", "Notifications in the basis"] + ["%d h" % T for T in SETTINGS] + [
        "SLA lateness (hours)"]
    _header(ws, 4, cols, numeric_from=3)
    ws.cell(row=3, column=4, value="Expected secondary cases a year if every notification had reached the unit this many "
            "hours after authorisation").font = Font(name="Calibri", size=9, italic=True, color="555555")
    hi = PatternFill("solid", fgColor="DCEBDF")
    r = 5
    for f in P.FEEDS:
        x = A["feeds"][f]
        vals = [W.FEED[f]["name"], W.FEED[f]["codes"][-1][0], x["n"]] + list(x["curve"]) + [x["setting"]]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.border = Border(bottom=THIN)
            if 4 <= j <= 10:
                c.number_format = "0.00"
                if SETTINGS[j - 4] == x["setting"]:
                    c.fill = hi
                    c.font = Font(name="Calibri", bold=True)
            if j == 11:
                c.font = Font(name="Calibri", bold=True)
                c.alignment = Alignment(horizontal="center")
        r += 1
    last = r - 1
    tr = r
    ws.cell(row=tr, column=1, value="Tolerance (CD-ST-07 section 2)").font = Font(name="Calibri", size=9, italic=True, color="666666")
    for j in range(4, 11):
        c = ws.cell(row=tr, column=j, value=P.TOLERANCE)
        c.font = Font(name="Calibri", size=9, italic=True, color="666666")
        c.number_format = "0.00"
    r += 2
    _note(ws, r, "Each feed is set at the longest allowed setting whose expected secondary cases a year stay at or under "
          "1.00 (shaded). The reference laboratory opened two investigations in three years, so even 24 hours costs it "
          "%.2f a year." % A["feeds"]["SERL"]["at"], 11, 150)
    r += 1
    _note(ws, r, "Expected secondary cases: for each notification, the child-hours of contact between the unit's call at "
          "authorisation and its call that many hours later (calls 07:00 to 21:00; the case leaves an hour after the call; "
          "booked sessions; the room's other attendees that day, from the centre's export), times the condition's "
          "secondary transmission (Transmission sheet); summed by feed and divided by three. Unrounded figures are carried "
          "through and rounded here to two decimals.", 11, 150)
    _widths(ws, [40, 15, 13, 8, 8, 8, 8, 8, 8, 8, 12])
    ch = LineChart()
    assert [f for f in P.FEEDS if A["feeds"][f]["next"] is None] == ["SERL"]
    ch.title = _title("Every feed but the reference laboratory goes over 1.00 a year at the next longer setting", 1200)
    ch.y_axis.title = _title("Expected secondary cases a year", 900, False)
    ch.x_axis.title = _title("SLA lateness (marker: the feed's setting)", 900, False)
    ch.height, ch.width = 10, 20
    data = Reference(ws, min_col=4, max_col=10, min_row=5, max_row=last)
    ch.add_data(data, from_rows=True, titles_from_data=False)
    for k, f in enumerate(P.FEEDS):
        sr = ch.series[k]
        sr.tx = SeriesLabel(v=SHORT[f])
        sr.smooth = False
        sr.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=LINE[f], w=22225))
        sr.marker = Marker(symbol="none")
        pt = DataPoint(idx=SETTINGS.index(A["feeds"][f]["setting"]))
        pt.marker = Marker(symbol="circle", size=8)
        pt.marker.graphicalProperties = GraphicalProperties(solidFill=LINE[f], ln=LineProperties(solidFill=LINE[f]))
        pt.graphicalProperties = GraphicalProperties(solidFill=LINE[f], ln=LineProperties(solidFill=LINE[f], w=22225))
        sr.dPt.append(pt)
    tol = Reference(ws, min_col=4, max_col=10, min_row=tr, max_row=tr)
    ch.add_data(tol, from_rows=True, titles_from_data=False)
    ch.series[-1].tx = SeriesLabel(v="Tolerance 1.00")
    ch.series[-1].smooth = False
    ch.series[-1].marker = Marker(symbol="none")
    ch.series[-1].graphicalProperties = GraphicalProperties(ln=LineProperties(prstDash="dash", solidFill="B3261E", w=19050))
    for sr in ch.series:
        sr.cat = AxDataSource(strRef=StrRef(f="'Settings'!$D$4:$J$4"))
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = 4
    ch.y_axis.majorUnit = 0.5
    ch.y_axis.majorGridlines = None
    ch.y_axis.number_format = "0.0"
    ch.legend.position = "b"
    ws.add_chart(ch, "A%d" % (r + 2))

    ws2 = wb.create_sheet("Transmission")
    _top(ws2, "Secondary transmission per 1,000 child-hours of contact: shigellosis %.2f, STEC infection %.2f, "
         "cryptosporidiosis %.2f" % tuple(1000 * A["rate"][c] for c in COND_ORDER),
         "The unit's childcare investigations from July 2021 that were closed at the 13 November 2026 extract "
         "(CD-ST-07 section 3).")
    _header(ws2, 4, ["Condition", "Closed investigations", "Secondary cases", "Child-hours of contact",
                     "Secondary cases per 1,000 child-hours", "Cases in the room follow-up"])
    r = 5
    for c in COND_ORDER:
        vals = [c, A["closed"][c], A["num"][c], round(A["den"][c], 2), round(1000 * A["rate"][c], 2), A["recorded"][c]]
        for j, v in enumerate(vals, start=1):
            cell = ws2.cell(row=r, column=j, value=v)
            cell.border = Border(bottom=THIN)
            if j == 4:
                cell.number_format = "#,##0.00"
            if j == 5:
                cell.number_format = "0.00"
                cell.font = Font(name="Calibri", bold=True)
        r += 1
    r += 1
    _note(ws2, r, "A secondary case is a room contact's confirmed or probable case (GL-CD-11 section 2) with onset no earlier "
          "than the index case's first day of infectious attendance plus the condition's shortest incubation period and no "
          "later than its last day of infectious attendance plus the longest (GL-CD-11 section 3: 12 hours to 7 days, 1 to "
          "10 days, 1 to 12 days). The follow-up also records cases ill before that window (the room's earlier and "
          "community cases) and after it (cases the secondary cases and later introductions brought); they are not the "
          "index case's secondary cases.", 6, 95)
    r += 1
    _note(ws2, r, "Child-hours of contact: on each day the index case attended from onset to the exclusion, the hours at "
          "the centre times the room's other children and educators that day (CD-ST-07 2.4(d) and 3.2).", 6, 95)
    _widths(ws2, [40, 13, 11, 14, 16, 14])

    ws3 = wb.create_sheet("Feed record")
    _top(ws3, "Reports authorised 1 October 2025 to 30 September 2026, and those first received more than 24 hours after "
         "authorisation", "Platform message log; a report counted once however many times it was received, at its first "
         "receipt.")
    _header(ws3, 4, ["Feed", "Reports", "First received more than 24 hours after authorisation"])
    r = 5
    for f in P.FEEDS:
        n, late = A["reports"][f]
        for j, v in enumerate([W.FEED[f]["name"], n, late], start=1):
            c = ws3.cell(row=r, column=j, value=v)
            c.border = Border(bottom=THIN)
            if j > 1:
                c.number_format = "#,##0"
        r += 1
    r += 1
    _note(ws3, r, "A report is one laboratory's authorisation of one test's result on one specimen: preliminary, final, "
          "corrected and amended reports each count. Re-sent copies (a disaster-recovery engine or a replay tool re-sending "
          "what had already gone, under new message ids and with the namespace on the accession number) are the same "
          "reports. Year by the New South Wales date of authorisation.", 3, 80)
    _widths(ws3, [40, 12, 22])

    ws4 = wb.create_sheet("Urgent cases")
    _top(ws4, "Shigellosis, STEC infection and cryptosporidiosis cases each feed was first to notify, 1 October 2025 to "
         "30 September 2026", "Platform case extract and message log.")
    _header(ws4, 4, ["Feed", "Cases first notified"])
    r = 5
    for f in P.FEEDS:
        for j, v in enumerate([W.FEED[f]["name"], A["cases"].get(f, 0)], start=1):
            c = ws4.cell(row=r, column=j, value=v)
            c.border = Border(bottom=THIN)
        r += 1
    r += 1
    _note(ws4, r, "A case's first notification is the first report the platform accepted for it, under its own number or "
          "a duplicate case merged into it; the year is the New South Wales date of that report. Confirmed and probable "
          "cases; merged duplicates are counted once, with their survivor.", 2, 60)
    _widths(ws4, [40, 18])

    ws5 = wb.create_sheet("Notes")
    for line in [
        "Prepared by Communicable Disease Surveillance, %s, 16 November 2026, for the NCP 4.2 feed register "
        "(go-live 1 December 2026)." % W.UNIT,
        "Sources: cc_investigation_register_2021-2026.xlsx, cc_contact_followup_2021-2026.csv, the three attendance exports "
        "and cc_room_bookings.csv; ncp_elr_message_log_2022-10_2026-10.csv, ncp_enteric_case_extract_2022-2026.csv and "
        "ncp_feed_register.xlsx, extracted 13 November 2026.",
        "Rules: CD-ST-07 version 1.0 (6 November 2026); GL-CD-11 version 3.1; NCP 4.2 release notes (14 October 2026).",
        "Times are New South Wales time unless marked UTC.",
    ]:
        ws5.append([line])
    ws5.column_dimensions["A"].width = 130
    for row in ws5.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.properties.creator = W.PEOPLE["manager"]
    wb.properties.lastModifiedBy = W.PEOPLE["manager"]
    wb.properties.title = "ELR feed SLA lateness basis 2026-27"
    wb.save(path)
    postprocess.scrub(path, W.PEOPLE["manager"], "2026-11-16")
    postprocess.finish_office(path, "2026-11-14T22:05:00Z", "2026-11-16T03:40:00Z", W.LHD)


def write(A, golden_dir):
    os.makedirs(golden_dir, exist_ok=True)
    write_csv(A, os.path.join(golden_dir, CSV))
    workbook(A, os.path.join(golden_dir, WORKBOOK))
