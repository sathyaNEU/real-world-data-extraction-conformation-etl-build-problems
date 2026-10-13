"""task232 assertions, run by make.py on every build: the three rules (R1 to R3c), the ladder and its margins, every
semantic fork row (SF1 to SF22), the convergences, the bins, Layer 1 neutrality, the asks' stops and separation, the
input gates and the realism sweep."""
import csv
import json
import math
import os
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta

import build as B
import estim as E
import golden as G
import params as P
import sim as S
import verify as V
import world as W

GOLD = {"CRBH": 12, "LPG": 2, "HDN": 12, "KES": 4, "HPS": 12, "SERL": 24}
DEFAULT = {"CRBH": 2, "LPG": 1, "HDN": 2, "KES": 2, "HPS": 8, "SERL": 24}
LOWER = {"CRBH": 12, "LPG": 1, "HDN": 4, "KES": 2, "HPS": 12, "SERL": 24}
BEFORE = {f: 24 for f in P.FEEDS}


def edge(x, step):
    y = x / step
    return abs(y - round(y - 0.5) - 0.5) * step


def margin(curve, tol=1.0):
    ok = [k for k, c in enumerate(curve) if c <= tol]
    s = max(ok)
    m = [tol / curve[s] - 1 if curve[s] > 0 else 9.0]
    if s + 1 < len(curve):
        m.append(curve[s + 1] / tol - 1)
    return min(m)


def run(target, golden_dir, world, plat, A, mainres, asks):
    results, fails = [], []

    def ok(label, cond, detail=""):
        results.append(label)
        if not cond:
            fails.append("%s: %s" % (label, detail))

    names = {W.FEED[f]["name"]: f for f in P.FEEDS}
    invs = world.invs
    closed = [i for i in invs if i.closed]

    # ------------------------------------------------------------------ input gates
    files = sorted(os.listdir(target))
    ok("gate: 10 or more files", len(files) >= 10, len(files))
    ok("gate: 3 or more formats", len({f.rsplit(".", 1)[1] for f in files}) >= 3, files)
    big = max(sum(1 for _ in open(os.path.join(target, f), encoding="utf-8")) - 1 for f in files if f.endswith(".csv"))
    ok("gate: a file of 25,000 rows or more", big >= 25000, big)
    meta = json.load(open(os.path.join(os.path.dirname(target), "metadata.json")))
    ok("gate: two or more distractors declared", len(meta["distractor_files"]) >= 2, meta["distractor_files"])
    ok("gate: distractors are shipped files", all(d in files for d in meta["distractor_files"]))
    ok("gate: no distractor is read by the verifier", not set(meta["distractor_files"]) & set(V.FILES.values()))
    ok("gate: no file name names the stump", not [f for f in files if re.search(r"distract|golden|answer|trap|clean|_v\d", f)])

    # ------------------------------------------------------------------ R1: the field's raw material ships, the quantity does not
    tr, tnum, tden = E.rates(invs, "window")
    ok("R1a: the verifier's secondary cases equal the truth's window count per condition",
       {c: A["num"][P.COND[c][0]] for c in P.COND_KEYS} == tnum, (A["num"], tnum))
    ok("R1a: the verifier's child-hours equal the truth's per condition",
       all(abs(A["den"][P.COND[c][0]] - tden[c]) < 1e-6 for c in P.COND_KEYS))
    heads = []
    for f in files:
        p = os.path.join(target, f)
        if f.endswith(".csv"):
            heads += next(csv.reader(open(p, encoding="utf-8")))
    ok("R1b: no shipped column holds a rate, a window or a secondary-case flag",
       not [h for h in heads if re.search(r"secondary|rate|window|attribut|transmission|generation|primary", h, re.I)],
       heads)
    fu = list(csv.DictReader(open(os.path.join(target, B.FILES["followup"]), encoding="utf-8")))
    ok("R1b: the follow-up's case status is only Confirmed, Probable or blank",
       set(r["case_status"] for r in fu) <= {"Confirmed", "Probable", ""})

    # ------------------------------------------------------------------ R3: the analyst test and the decisive step
    def st_truth(rung):
        R = E.rates(invs, rung)[0]
        C = E.costs(invs, R)
        return {f: E.setting(C[f]) for f in P.FEEDS}, C
    g_truth, Cg = st_truth("window")
    d_truth, Cd = st_truth("all")
    l_truth, Cl = st_truth("lower")
    b_truth, Cb = st_truth("before")
    ok("R3a: the default path (every case ill after the index) on the clean truth files Ends on", l_truth == LOWER, l_truth)
    ok("R3a: the default path from the shipped files files Ends on", A["rungs"]["lower"]["settings"] == LOWER,
       A["rungs"]["lower"]["settings"])
    ok("R3a: R0 (every follow-up case) on the clean truth and from the shipped files",
       d_truth == DEFAULT and A["rungs"]["all"]["settings"] == DEFAULT, (d_truth, A["rungs"]["all"]["settings"]))
    ok("R3b: the attribution alone files the golden (truth)", g_truth == GOLD, g_truth)
    ok("R3b: the attribution alone files the golden (shipped files)", {f: A["feeds"][f]["setting"] for f in P.FEEDS} == GOLD)
    ok("R3b: R1's single correction (drop cases before the index) leaves three feeds wrong",
       A["rungs"]["lower"]["settings"] == LOWER and sum(LOWER[f] != GOLD[f] for f in P.FEEDS) == 3,
       A["rungs"]["lower"]["settings"])
    ok("R3b: R2 (cases before the exclusion) sets every feed at 24", A["rungs"]["before"]["settings"] == BEFORE)
    ok("ladder: R0, R1, R2 and the golden file four different settings vectors",
       len({tuple(sorted(x.items())) for x in (DEFAULT, LOWER, BEFORE, GOLD)}) == 4)
    ok("ladder: R0 differs from the golden on five feeds", sum(DEFAULT[f] != GOLD[f] for f in P.FEEDS) == 5)
    ok("ladder: every rung moves all three rates",
       all(abs(A["rungs"][r]["rate"][c] - A["rate"][c]) * 1000 > 0.3 for r in ("all", "lower", "before")
           for c in A["rate"]))
    gm = min(margin(Cg[f]) for f in P.FEEDS)
    dm = min(margin(Cd[f]) for f in P.FEEDS)
    ok("ladder: golden margins at least 20 per cent", gm >= 0.20, gm)
    ok("ladder: R0 margins at least 15 per cent", dm >= 0.15, dm)
    wrong = [f for f in P.FEEDS if LOWER[f] != GOLD[f]]
    lm = min(margin(Cl[f]) for f in wrong)
    ok("ladder: the default path's three wrong feeds sit at least 5 per cent from the tolerance", lm >= 0.05, lm)
    reg_head = list(V.rows_xlsx(os.path.join(target, B.FILES["register"]), "Register")[0].keys())
    ok("R3c: the register carries no case count, rate or setting",
       not [h for h in reg_head if re.search(r"case(s)? (count|in)|rate|setting|secondary|attack", str(h), re.I)], reg_head)
    std = V.text_docx(os.path.join(target, B.FILES["standard"]))
    rel = V.text_pdf(os.path.join(target, B.FILES["release"]))
    ok("R3c: the release notes give no default lateness", not re.search(r"default", rel, re.I))
    ok("R3c: no shipped document states a setting for a feed",
       not re.search(r"(Larkspur|Halden|Kestrel|Highfield|reference laboratory|base hospital)[^.]{0,60}\b(1|2|4|6|8|12|24) hours",
                     std + rel, re.I))
    ok("R3c: the standard does not say which room cases are secondary",
       not re.search(r"incubation|window|onset", std, re.I))
    gl = re.sub(r"\s+", " ", V.text_pdf(os.path.join(target, B.FILES["guideline"])))
    ok("R3c: the guideline does not say how a case is attributed to the index",
       not re.search(r"within (the|one) incubation|incubation period after|after (the|its) exposure|onset (falls|within)",
                     gl, re.I))

    # ------------------------------------------------------------------ semantic fork rows
    ok("SF1: six feeds; the regional laboratory's two codes are one feed in the register",
       len(set(V.feed_of_codes(target).values())) == 6 and len(V.feed_of_codes(target)) == 7)
    reg = V.rows_xlsx(os.path.join(target, B.FILES["register"]), "Register")
    ok("SF2: one opening notification per investigation", len(reg) == len({r["Investigation"] for r in reg}) == len(invs))
    def basis_by(fn):
        return Counter(i.feed for i in invs if P.BASIS_FROM <= fn(i) <= P.BASIS_TO)
    bset = basis_by(lambda i: i.auth.date())
    ok("SF3: the basis is the same 80 by authorisation, receipt, call, specimen or onset date, NSW or UTC",
       all(basis_by(fn) == bset for fn in (lambda i: i.received.date(), lambda i: i.call.date(), lambda i: i.specimen,
                                           lambda i: i.onset, lambda i: i.auth.astimezone(W.UTC).date(),
                                           lambda i: i.received.astimezone(W.UTC).date())) and sum(bset.values()) == 80,
       dict(bset))
    ok("SF4: a year is the three years' total divided by three", all(
        abs(A["feeds"][f]["curve"][P.SETTINGS.index(A["feeds"][f]["setting"])] - A["feeds"][f]["at"]) < 1e-12 for f in P.FEEDS))
    Rc = E.rates(invs, "window", confirmed_only=True)[0]
    Cc = E.costs(invs, Rc)
    ok("SF5: confirmed-only cases move the call (ruled out by CD-ST-07 3.3)",
       {f: E.setting(Cc[f]) for f in P.FEEDS} != GOLD and "confirmed and probable cases" in std)
    num_o = defaultdict(int)
    den_o = defaultdict(float)
    for i in invs:
        den_o[i.cond] += E.contact_hours(i)
        num_o[i.cond] += sum(1 for x in i.cases if E.keep(i, "window", x["onset"]) and x["onset"] <= P.EXTRACT)
    Ro = {c: num_o[c] / den_o[c] for c in P.COND_KEYS}
    Co = E.costs(invs, Ro)
    ok("SF6: open investigations included give the same call", {f: E.setting(Co[f]) for f in P.FEEDS} == GOLD)
    Rm = {}
    for c in P.COND_KEYS:
        rs = [sum(1 for x in i.cases if E.keep(i, "window", x["onset"])) / E.contact_hours(i) for i in closed if i.cond == c]
        Rm[c] = sum(rs) / len(rs)
    Cm = E.costs(invs, Rm)
    ok("SF7: the mean of investigation rates gives the same call", {f: E.setting(Cm[f]) for f in P.FEEDS} == GOLD)

    def ch_pair(i):
        tot = 0.0
        for d in i.inf_days:
            a, b = i.att[(i.index["pid"], d)]
            for p in i.roster.people:
                if p is not i.index and (p["pid"], d) in i.att:
                    x, y = i.att[(p["pid"], d)]
                    tot += max(0.0, min(b, y) - max(a, x))
        return tot

    def extra_pair(i, T):
        c0, cT = W.call_time(i.auth), W.call_time(W.add_hours(i.auth, T))
        s0 = i.index["start"]
        tot, d = 0.0, c0.date()
        while d <= cT.date():
            if W.centre_open(i.centre, d) and d.weekday() in i.index["days"]:
                e0 = s0 + E.held(c0, d, s0, i.index["end"])
                e1 = s0 + E.held(cT, d, s0, i.index["end"])
                for p in i.roster.people:
                    if p is not i.index and (p["pid"], d) in i.att:
                        x, y = i.att[(p["pid"], d)]
                        tot += max(0.0, min(e1, y) - max(s0, x)) - max(0.0, min(e0, y) - max(s0, x))
            d += timedelta(days=1)
        return tot
    Rp = {c: tnum[c] / sum(ch_pair(i) for i in closed if i.cond == c) for c in P.COND_KEYS}
    Cp = {f: [sum(Rp[i.cond] * extra_pair(i, T) for i in E.basis(invs) if i.feed == f) / 3.0 for T in P.SETTINGS]
          for f in P.FEEDS}
    ok("SF8: each pair's overlapping hours (rate and exposure) give the same call",
       {f: E.setting(Cp[f]) for f in P.FEEDS} == GOLD, {f: E.setting(Cp[f]) for f in P.FEEDS})
    ok("SF9: infectious attendance starts on the onset date for every index",
       all(i.first_inf >= i.onset and i.inf_days[0] == i.first_inf for i in invs))
    ok("SF10: the counterfactual is anchored at authorisation (CD-ST-07 2.4)", "at authorisation" in std)
    old = P.LEAVE_AFTER
    P.LEAVE_AFTER = 0.0
    try:
        C0 = E.costs(invs, tr)
    finally:
        P.LEAVE_AFTER = old
    ok("SF11: leaving at the call gives the same call; a whole session is ruled out by CD-ST-07 2.4(b)",
       {f: E.setting(C0[f]) for f in P.FEEDS} == GOLD and "leaves one hour after the call" in std)
    ok("SF12: the counterfactual uses booked sessions (CD-ST-07 2.4(c))", "booked or rostered" in std)
    oldo = E.others

    def others_booked(i, d):
        return sum(1 for p in i.roster.people if p is not i.index and d.weekday() in p["days"]) if W.centre_open(i.centre, d) else 0
    E.others = others_booked
    try:
        Cb2 = {f: [sum(tr[i.cond] * E.extra(i, T) for i in E.basis(invs) if i.feed == f) / 3.0 for T in P.SETTINGS]
               for f in P.FEEDS}
    finally:
        E.others = oldo
    ok("SF13: the room's bookings for the others count give the same call", {f: E.setting(Cb2[f]) for f in P.FEEDS} == GOLD)
    ok("SF14: no expected-case figure within 0.005 of 1.00",
       min(abs(c - 1.0) for f in P.FEEDS for c in A["feeds"][f]["curve"]) > 0.005)
    ok("SF15: the reference laboratory at 24 hours has no longer setting", A["feeds"]["SERL"]["next"] is None)
    rr = {c: round(1000 * v, 2) / 1000 for c, v in tr.items()}
    Cr = E.costs(invs, rr)
    def figs(C, st):
        out = []
        for f in P.FEEDS:
            k = P.SETTINGS.index(st[f])
            out.append(round(C[f][k], 2))
            if k + 1 < 7:
                out.append(round(C[f][k + 1], 2))
        return out
    ok("SF16: rounding the rates first gives the same figures", figs(Cr, GOLD) == figs(Cg, GOLD))

    def extra_wall(i, T):
        c0 = W.call_time(i.auth)
        t = i.auth.replace(tzinfo=None) + timedelta(hours=T)
        cT = W.call_time(t.replace(tzinfo=W.NSW))
        tot = 0.0
        d = c0.date()
        while d <= cT.date():
            if W.centre_open(i.centre, d) and d.weekday() in i.index["days"]:
                h = E.held(cT, d, i.index["start"], i.index["end"]) - E.held(c0, d, i.index["start"], i.index["end"])
                if h > 0:
                    tot += h * E.others(i, d)
            d += timedelta(days=1)
        return tot
    Cw = {f: [sum(tr[i.cond] * extra_wall(i, T) for i in E.basis(invs) if i.feed == f) / 3.0 for T in P.SETTINGS]
          for f in P.FEEDS}
    ok("SF17: wall-clock hours give the same figures", figs(Cw, GOLD) == figs(Cg, GOLD))
    zone_hits = 0
    for i in closed:
        (l0, l1), (h0, h1) = world.zones(i)
        zone_hits += sum(1 for x in i.cases if l0 <= x["onset"] <= l1 or h0 <= x["onset"] <= h1)
    ok("SF18: no case onset in an edge zone (anchors on attendance or on onset and exclusion agree)", zone_hits == 0,
       zone_hits)
    diff = 0
    for i in closed:
        lo, hi = P.COND[i.cond][1], P.COND[i.cond][2]
        for x in i.cases:
            days = [d for d in i.inf_days if (x["pid"], d) in i.att]
            own = bool(days) and (x["onset"] - days[0]).days >= lo and (x["onset"] - days[-1]).days <= hi
            diff += own != E.keep(i, "window", x["onset"])
    ok("SF19: each contact's own shared days keep the same cases as the room's window", diff == 0, diff)
    gap = sum(1 for i in closed for x in i.cases if i.onset <= x["onset"] and not E.keep(i, "lower", x["onset"]))
    ok("SF20: no case between the index's onset and its first infectious attendance plus the shortest incubation",
       gap == 0, gap)
    ok("SF21: the tolerance is 1.00 in the standard", "do not exceed 1.00" in std)
    ok("SF22: the guideline defines a secondary case as a room contact's case acquired from the index case",
       "Secondary case. A confirmed or probable case in a room contact who acquired the infection from the index case." in gl)

    # ------------------------------------------------------------------ C1 conventions on the clock
    ok("C1: no booked or rostered index session starts at 07:00",
       all(i.index["start"] != P.CALL_FROM for i in invs))
    ok("C1: no authorisation on the quarter hour (no call ties with a booked arrival or 07:00 or 21:00)",
       all(i.auth.minute % 15 != 0 for i in invs))
    ok("C1: no two investigations share a room within their exports",
       not [1 for a in invs for b in invs if a is not b and a.centre is b.centre and a.room is b.room and
            not (b.ex_from > a.ex_to or a.ex_from > b.ex_to)])

    # ------------------------------------------------------------------ bins
    gf = []
    for f in P.FEEDS:
        gf.append(A["feeds"][f]["at"])
        if A["feeds"][f]["next"] is not None:
            gf.append(A["feeds"][f]["next"])
    ok("bins: every expected-case figure at least 0.0012 from its rounding edge", min(edge(x, 0.01) for x in gf) >= 0.0012,
       min(edge(x, 0.01) for x in gf))
    ok("bins: every rate at least 0.0015 from its rounding edge",
       min(edge(1000 * v, 0.01) for v in A["rate"].values()) >= 0.0015)
    ok("bins: no graded figure is a whole or half number",
       not [x for x in gf + [1000 * v for v in A["rate"].values()] if round(x, 2) * 100 % 50 == 0])

    # ------------------------------------------------------------------ R2: Layer 1 items, mishandled, leave the call
    att_a = list(csv.DictReader(open(os.path.join(target, B.FILES["att_a"]), encoding="utf-8")))
    ok("R2a: the sign-in log's dates parse one way (DD-Mon-YYYY)",
       all(re.fullmatch(r"\d{2}-[A-Z][a-z]{2}-\d{4}", r["date"]) for r in att_a))
    roll = V.rows_xlsx(os.path.join(target, B.FILES["att_b"]), "Roll")
    pm = [r for r in roll if str(r["Out"]).endswith("PM")]
    bad = 0
    for r in pm[:200]:
        try:
            datetime.strptime(r["Out"], "%H:%M")
        except ValueError:
            bad += 1
    ok("R2b: a 24-hour read of the daily roll's 12-hour times fails loudly", bad == min(200, len(pm)) and pm)
    vin, vnames = V.load_investigations(target)
    truth_hours = {}
    for i in invs:
        truth_hours[i.iid] = sum((b - a) for (pid, d), (a, b) in i.att.items())
    ship_hours = {v["iid"]: sum(b - a for day in v["att"].values() for (a, b) in day.values()) for v in vin}
    ok("R2c: every layout parses to the attendance the centres recorded",
       all(abs(truth_hours[k] - ship_hours[k]) < 1e-6 for k in truth_hours), "")
    ok("R2d: every index person number is in its own centre's export",
       all(any(v["index"] in day for day in v["att"].values()) for v in vin))
    ok("R2e: casual educators appear in the attendance only, never in the bookings",
       all(not any(p["pid"] in v["book"] for p in i.roster.people if not p["days"]) for i, v in
           zip(sorted(invs, key=lambda x: x.iid), sorted(vin, key=lambda x: x["iid"]))))
    appro = [r for r in fu if r["onset_approximate"] == "Y"]
    vmap = {v["iid"]: v for v in vin}

    F = V.facts(target)
    steady = 0
    for r in appro:
        v = vmap[r["investigation_id"]]
        lo, hi = F["incubation"][v["cond"]]
        o = date.fromisoformat(r["onset_date"])
        member = {((o + timedelta(days=k)) - v["inf_from"]).days >= lo and ((o + timedelta(days=k)) - v["inf_to"]).days <= hi
                  for k in range(-2, 3)}
        steady += len(member) == 1
    ok("R2f: every approximate onset keeps its place in the window two days either side", steady == len(appro) and appro,
       (steady, len(appro)))
    ok("R2g: the call reads none of the platform's files (message log, case extract, incidents, code set)",
       not {B.FILES[k] for k in ("log", "extract", "incidents", "codes")} & set(
           [V.FILES[k] for k in ("register", "followup", "att_a", "att_b", "att_c", "bookings", "feeds", "standard",
                                 "guideline", "release")]))

    # ------------------------------------------------------------------ the asks
    msgs = asks["msgs"]
    ans_r = {names[n]: v for n, v in asks["reports"].items()}
    ans_c = {names[n]: v for n, v in asks["cases"].items()}
    ok("ask A: the generator's report count agrees with the verifier",
       ans_r == A["reports"], (ans_r, A["reports"]))
    stops = {
        "every message": V.ask_reports(msgs, lambda m: m["msg_control_id"]),
        "accession, test and status": V.ask_reports(msgs, lambda m: (m["_feed"], m["accession_no"], m["test_code"],
                                                                     m["result_status"])),
        "accession and test": V.ask_reports(msgs, lambda m: (m["_feed"], m["accession_no"], m["test_code"])),
        "prefix stripped, accession, test and status": V.ask_reports(msgs, lambda m: (m["_feed"], m["accession_no"].split("^")[0],
                                                                                      m["test_code"], m["result_status"])),
        "accession, test, status and authorisation": V.ask_reports(msgs, lambda m: (m["_feed"], m["accession_no"], m["test_code"],
                                                                                    m["result_status"], m["authorised_at"])),
        "every namespaced accession dropped": V.ask_reports([m for m in msgs if "^" not in m["accession_no"]],
                                                            lambda m: m["msg_control_id"]),
        "namespaced accessions dropped, accession, test and status": V.ask_reports(
            [m for m in msgs if "^" not in m["accession_no"]],
            lambda m: (m["_feed"], m["accession_no"], m["test_code"], m["result_status"])),
    }
    ok("ask A: keyed on the accession without its namespace, test, status and authorisation, the answer is the same",
       {names[n]: v for n, v in V.ask_reports(msgs, lambda m: (m["_feed"], m["accession_no"].split("^")[0], m["test_code"],
                                                              m["result_status"], m["authorised_at"])).items()} == ans_r)
    for k, st in stops.items():
        st = {names[n]: v for n, v in st.items()}
        ok("ask A stop (%s) misses every feed's reports" % k, all(st[f][0] != ans_r[f][0] for f in P.FEEDS), st)
        ok("ask A stop (%s) misses every feed's late reports" % k, all(st[f][1] != ans_r[f][1] for f in P.FEEDS), st)
    for mode in ("every", "live"):
        st = {names[n]: v for n, v in V.ask_cases(target, msgs, mode=mode).items()}
        ok("ask B stop (%s case numbers) misses every feed" % mode, all(st.get(f, 0) != ans_c.get(f, 0) for f in P.FEEDS),
           st)
    ext = list(csv.DictReader(open(os.path.join(target, B.FILES["extract"]), encoding="utf-8")))
    lab = Counter()
    for e in ext:
        if e["condition"] in V.URGENT and e["case_status"] in ("Confirmed", "Probable") and e["notifying_facility"] and \
                "2025-10-01" <= e["notified_date"] <= "2026-09-30":
            lab[names[V.feed_of_codes(target)[e["notifying_facility"]]]] += 1
    ok("ask B stop (the extract's stored notifying facility) misses every feed",
       all(lab.get(f, 0) != ans_c.get(f, 0) for f in P.FEEDS), dict(lab))
    inv_cases = {id(i.case) for i in invs} | {id(x["case"]) for i in invs for x in i.cases}
    ok("separation: no re-sent copy or amended report of an investigation case",
       not [x for x in plat.reports if x["kind"] != "original" and id(x["case"]) in inv_cases])
    ok("separation: no merged or second-episode case among the investigation cases",
       not [1 for a, b in plat.merge_log + plat.reinfections if id(a) in inv_cases or id(b) in inv_cases])
    ok("separation: no investigation report delivered by a disaster-recovery engine",
       not [x for x in plat.reports if x.get("via_dr") and id(x["case"]) in inv_cases])
    ok("ask devices: every feed has re-sent copies or amended reports in the year",
       all(any(x["feed"] == f and x["kind"] in ("copy", "amended") and P.YEAR_FROM <= x["auth"].astimezone(W.NSW).date()
               <= P.YEAR_TO for x in plat.reports) for f in P.FEEDS))
    ok("ask B devices: six merges and three second episodes", len(plat.merge_log) == 6 and len(plat.reinfections) == 3)

    # ------------------------------------------------------------------ the goldens agree with the answers
    gcsv = list(csv.DictReader(open(os.path.join(golden_dir, G.CSV), encoding="utf-8")))
    ok("golden: the CSV carries every feed's setting and figures",
       [(r["sla_lateness_hours"], r["expected_secondary_cases_a_year_at_setting"],
         r["expected_secondary_cases_a_year_at_next_longer_setting"]) for r in gcsv] ==
       [(str(A["feeds"][f]["setting"]), "%.2f" % A["feeds"][f]["at"],
         "%.2f" % A["feeds"][f]["next"] if A["feeds"][f]["next"] is not None else "") for f in P.FEEDS])

    import openpyxl
    gwb = openpyxl.load_workbook(os.path.join(golden_dir, G.WORKBOOK))
    st = gwb["Settings"]
    srow = {st.cell(row=r, column=2).value: [st.cell(row=r, column=c).value for c in range(3, 12)] for r in range(5, 11)}
    ok("golden: the workbook's settings and expected cases are the CSV's",
       all(srow[W.FEED[f]["codes"][-1][0]][-1] == A["feeds"][f]["setting"]
           and "%.2f" % srow[W.FEED[f]["codes"][-1][0]][1 + P.SETTINGS.index(A["feeds"][f]["setting"])] ==
           "%.2f" % A["feeds"][f]["at"] for f in P.FEEDS))
    tw = gwb["Transmission"]
    ok("golden: the workbook's rates are the answers",
       ["%.2f" % tw.cell(row=r, column=5).value for r in range(5, 8)] ==
       ["%.2f" % (1000 * A["rate"][c]) for c in G.COND_ORDER])
    fw, uw = gwb["Feed record"], gwb["Urgent cases"]
    ok("golden: the workbook's report counts and case counts are the answers",
       [(fw.cell(row=r, column=2).value, fw.cell(row=r, column=3).value, uw.cell(row=r, column=2).value)
        for r in range(5, 11)] == [(A["reports"][f][0], A["reports"][f][1], A["cases"].get(f, 0)) for f in P.FEEDS])

    # ------------------------------------------------------------------ realism sweep
    text_files = [f for f in files if f.endswith((".csv", ".json"))]
    dash = [f for f in text_files if re.search("[–—]", open(os.path.join(target, f), encoding="utf-8").read())]
    ok("realism: no em or en dash in any text file", not dash, dash)
    ok("realism: no em or en dash in the documents", not re.search("[–—]", std + rel))
    card = json.load(open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(target))), ".claude", "skills",
                                       "fingerprint", "cards", "task232.json")))
    ok("realism: the people are the card's", set(W.PEOPLE.values()) == set(card["world"]["people"]))
    rep = Counter(re.sub(r"[0-9.:]+", "#", r[6]) for r in csv.reader(open(os.path.join(target, B.FILES["incidents"]),
                                                                         encoding="utf-8")))
    ok("realism: no incident summary recurs more than five times, figures and times aside", max(rep.values()) <= 5,
       rep.most_common(2))
    counts = [A["reports"][f][0] for f in P.FEEDS]
    ok("realism: no two feeds share a report count", len(set(counts)) == len(counts))
    log_last = max(m["_recv"] for m in msgs).astimezone(W.NSW).date()
    ok("realism: no message after the extract's end", log_last <= date(2026, 10, 31), log_last)
    ok("realism: no investigation closed after the extract", all(i.closed_on <= P.EXTRACT for i in invs if i.closed))
    ok("realism: the log carries 25,000 or more messages", len(msgs) >= 25000, len(msgs))

    if fails:
        raise AssertionError("%d check(s) failed:\n  " % len(fails) + "\n  ".join(fails))
    return results
