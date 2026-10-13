"""task232 documents: the unit's feed SLA standard (CD-ST-07), its childcare enteric response guideline, the platform's
4.2 release notes on per-feed SLA alerts, the feed register, the investigation register workbook, the centres' daily
roll workbook and the data dictionary."""
from datetime import date, datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess
import world as W

COMPANY = W.LHD


# ------------------------------------------------------------------------------------------------ helpers
def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=13, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=14, leading=18, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=14, spaceBefore=8,
                        spaceAfter=4)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10, textColor=colors.HexColor("#444444"))
    cell = ParagraphStyle("cell", parent=body, fontSize=8.5, leading=10.5, spaceAfter=0)
    item = ParagraphStyle("item", parent=body, leftIndent=14, spaceAfter=3)
    return dict(body=body, h1=h1, h2=h2, small=small, cell=cell, item=item)


def _footer(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(18 * mm, 10 * mm, text)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _pdf(path, story, title, author, footer, stamp):
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title=title, author=author, creator=author, invariant=1)
    doc.build(story, onFirstPage=_footer(footer), onLaterPages=_footer(footer))
    postprocess.scrub(path, author, stamp)


def _table(rows, widths, size=8.5):
    t = Table(rows, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", size), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", size),
                           ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E2E9E4")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


HEAD = Font(name="Calibri", bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="2F5E4E")


def _sheet(ws, header, rows, widths=None, wrap=False, formats=None):
    ws.append(header)
    for c in ws[1]:
        c.font = HEAD
        c.fill = FILL
        c.alignment = Alignment(vertical="top", wrap_text=True)
    for r in rows:
        ws.append(r)
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    if formats:
        for col, fmt in formats.items():
            for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                for c in row:
                    c.number_format = fmt
    if wrap:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


def _save(wb, path, who, created, modified):
    wb.properties.creator = who
    wb.properties.lastModifiedBy = who
    wb.save(path)
    postprocess.scrub(path, who, modified[:10])
    postprocess.finish_office(path, created, modified, COMPANY)


def _naive(t):
    """an aware time as New South Wales wall-clock, for a workbook cell"""
    return t.astimezone(W.NSW).replace(tzinfo=None)


# ------------------------------------------------------------------------------------------------ the standard (docx)
STANDARD = [
    ("1 Purpose", [
        "From 1 December 2026 release 4.2 of the Notifiable Conditions Platform (NCP) raises an SLA alert for each "
        "laboratory's electronic notification feed: a report that reaches the platform later than the feed's SLA lateness "
        "after the laboratory authorised it pages the unit's on-call public health officer, who acts on the report at once "
        "and takes the delay up with the laboratory. This standard sets each feed's SLA lateness.",
        "The unit prices a late notification by what it costs first: the time an infectious child or educator stays in "
        "their room at a long day care centre before the unit's exclusion call."]),
    ("2 Setting a feed's SLA lateness", [
        "2.1 A feed's SLA lateness is the longest of the platform's allowed settings (1, 2, 4, 6, 8, 12 and 24 hours) at "
        "which the feed's expected secondary cases a year do not exceed 1.00. A feed whose expected secondary cases exceed "
        "1.00 even at 1 hour is set at 1 hour.",
        "2.2 A feed's expected secondary cases a year at a setting are the expected secondary cases from the feed's urgent "
        "childcare notifications in the basis period, had each reached the unit that many hours after the laboratory "
        "authorised it rather than at authorisation, divided by three.",
        "2.3 The basis period is 1 October 2023 to 30 September 2026. A feed's urgent childcare notifications are the "
        "notifications that opened the unit's childcare investigations, as the childcare investigation register records "
        "them, with a result authorised in the basis period.",
        "2.4 For one notification at one setting, the expected secondary cases are the child-hours of contact the notified "
        "case would have added in their room, multiplied by the condition's secondary transmission per child-hour of contact "
        "(section 3). The child-hours added are those between the unit's call had the notification reached the unit at "
        "authorisation and the unit's call had it reached the unit that many hours later, where:"]),
]
STANDARD_ITEMS = [
    "(a) the unit calls the family and the centre at the first moment in its calling hours, 07:00 to 21:00 New South Wales "
    "time every day, at or after the notification reaches it;",
    "(b) a case who is at the centre when the unit calls leaves one hour after the call, a case who is not yet at the "
    "centre does not attend that day, and no case attends a session that starts after the call;",
    "(c) the case is taken to attend every session they are booked or rostered for, at the booked or rostered times; and",
    "(d) child-hours of contact on a day are the hours the case is at the centre that day multiplied by the number of other "
    "children and educators who attended the case's room that day.",
]
STANDARD_REST = [
    ("3 Secondary transmission", [
        "3.1 A condition's secondary transmission per child-hour of contact is the number of secondary cases among room "
        "contacts in the unit's closed childcare investigations of the condition, divided by the child-hours of contact "
        "during those investigations' infectious attendance.",
        "3.2 An index case's infectious attendance is their attendance at the centre from the start of their infectious "
        "period to their exclusion, from the attendance records the centre supplied to the investigation. Child-hours of "
        "contact are counted on each day of it as in 2.4(d), from the hours the index case was at the centre.",
        "3.3 Cases are confirmed and probable cases under the case definitions in the unit's guideline for enteric infections "
        "in early childhood education and care."]),
    ("4 Reporting", [
        "4.1 The surveillance team reports each feed's setting with its expected secondary cases a year at that setting and "
        "at the next longer allowed setting, to two decimal places, and each condition's secondary transmission per 1,000 "
        "child-hours of contact, to two decimal places. Figures are rounded only where they are reported.",
        "4.2 The settings go to the NCP feed register for the 4.2 release and are reviewed each November on the three years "
        "to the preceding 30 September."]),
]


def _docx_footer(doc, text):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text + "  |  Page ")
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    run = p.add_run()
    run.font.size = Pt(7.5)
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)


def standard(path):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(5)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.3)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
    _docx_footer(doc, "CD-ST-07 Laboratory feed SLA lateness, version 1.0  |  %s" % W.UNIT)
    h = doc.add_paragraph()
    r = h.add_run("CD-ST-07  Laboratory feed SLA lateness")
    r.bold = True
    r.font.size = Pt(15)
    meta = doc.add_paragraph()
    mr = meta.add_run("%s, communicable disease standard\nVersion 1.0, approved 6 November 2026 by %s, Director, Public "
                      "Health Unit\nOwner: %s, Manager, Communicable Disease Surveillance\nApplies from the NCP 4.2 release, "
                      "1 December 2026" % (W.UNIT, W.PEOPLE["director"], W.PEOPLE["manager"]))
    mr.font.size = Pt(9)
    mr.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    def heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        rr = p.add_run(text)
        rr.bold = True
        rr.font.size = Pt(11)

    for head, paras in STANDARD:
        heading(head)
        for t in paras:
            doc.add_paragraph(t)
    for t in STANDARD_ITEMS:
        p = doc.add_paragraph(t)
        p.paragraph_format.left_indent = Cm(0.8)
    for head, paras in STANDARD_REST:
        heading(head)
        for t in paras:
            doc.add_paragraph(t)
    doc.core_properties.author = W.PEOPLE["manager"]
    doc.core_properties.title = "CD-ST-07 Laboratory feed SLA lateness"
    doc.save(path)
    postprocess.scrub(path, W.PEOPLE["manager"], "2026-11-06")
    postprocess.finish_office(path, "2026-09-22T03:10:00Z", "2026-11-06T04:42:00Z", COMPANY)


# ------------------------------------------------------------------------------------------------ the guideline (pdf)
def guideline(path):
    S = _styles()
    st = [Paragraph("Enteric infections in early childhood education and care: case and room response", S["h1"]),
          Paragraph("%s. Guideline GL-CD-11, version 3.1, May 2024. Owner: Manager, Communicable Disease Surveillance. "
                    "Replaces version 3.0 (June 2022)." % W.UNIT, S["small"]), Spacer(1, 6)]
    secs = [
        ("1 Scope", [
            "This guideline covers a notified case of shigellosis, Shiga toxin-producing E. coli (STEC) infection or "
            "cryptosporidiosis in a child who attends, or an educator who works at, a long day care centre and was there "
            "while infectious. The unit opens a childcare investigation for each such case (the index case)."]),
        ("2 Case definitions for a room investigation", [
            "<b>Confirmed case.</b> A room contact with a laboratory result positive for the condition.",
            "<b>Probable case.</b> A room contact with diarrhoea who has an epidemiological link to a confirmed case in the "
            "room and was not tested, or whose specimen could not be tested.",
            "<b>Secondary case.</b> A confirmed or probable case in a room contact who acquired the infection from the index "
            "case.",
            "A room contact is a child or educator who attended the index case's room on a day the index case attended "
            "while infectious."]),
        ("3 The conditions", []),
    ]
    for head, paras in secs:
        st.append(Paragraph(head, S["h2"]))
        for p in paras:
            st.append(Paragraph(p, S["body"]))
    c = S["cell"]
    st.append(_table([
        ["Condition", "Incubation period", "Infectious period", "Exclusion of a case"],
        [Paragraph("Shigellosis", c), Paragraph("12 hours to 7 days", c),
         Paragraph("From onset of symptoms while the organism is in the faeces; up to four weeks without treatment.", c),
         Paragraph("Until 48 hours after the last loose stool and one negative stool specimen, taken at least 48 hours "
                   "after any antibiotic.", c)],
        [Paragraph("STEC infection", c), Paragraph("1 to 10 days", c),
         Paragraph("From onset of symptoms while the organism is shed; often three weeks in young children.", c),
         Paragraph("Until two consecutive negative stool specimens taken at least 24 hours apart.", c)],
        [Paragraph("Cryptosporidiosis", c), Paragraph("1 to 12 days", c),
         Paragraph("From onset of symptoms while oocysts are shed; can continue for weeks after symptoms stop.", c),
         Paragraph("Until 24 hours after the last loose stool; no swimming for two weeks after.", c)],
    ], [30 * mm, 30 * mm, 56 * mm, 58 * mm]))
    st.append(Spacer(1, 4))
    more = [
        ("4 On notification", [
            "4.1 The duty officer calls the family and the centre and excludes the case from the centre as in section 3.",
            "4.2 The officer asks the centre for the room's attendance from 21 days before the case's onset to 14 days after "
            "the exclusion, as its management software exports it, with the room's bookings and the educators' rosters.",
            "4.3 From the attendance the officer lists the room contacts and asks the centre and the families about "
            "diarrhoea in the room from 21 days before the case's onset. Room contacts with diarrhoea are asked to give a "
            "stool specimen.",
            "4.4 The officer follows the room until 14 days (shigellosis), 20 days (STEC infection) or 24 days "
            "(cryptosporidiosis) after the exclusion, records each contact's illness, onset date, specimen and result and "
            "case status in the contact follow-up, and enters confirmed and probable cases on NCP.",
            "4.5 The investigation is closed when the follow-up period has ended and outstanding results are in."]),
        ("5 Records", [
            "The childcare investigation register (one row per investigation), the contact follow-up (one row per room "
            "contact) and the centre's exports as supplied are kept in the unit's investigation folder for seven years."]),
    ]
    for head, paras in more:
        st.append(Paragraph(head, S["h2"]))
        for p in paras:
            st.append(Paragraph(p, S["body"]))
    _pdf(path, st, "GL-CD-11 Enteric infections in early childhood education and care, version 3.1", W.PEOPLE["manager"],
         "GL-CD-11 version 3.1  |  %s" % W.UNIT, "2024-05-20")


# ------------------------------------------------------------------------------------------------ release notes (pdf)
def release_notes(path):
    S = _styles()
    st = [Paragraph("Notifiable Conditions Platform release 4.2: per-feed SLA alerts", S["h1"]),
          Paragraph("Release notes for public health units (extract). State notifiable conditions program, platform "
                    "services. Issued 14 October 2026. Go-live 1 December 2026.", S["small"]), Spacer(1, 6)]
    for head, paras in [
        ("What changes", [
            "Release 4.2 adds an SLA alert to each electronic laboratory notification (ELR) feed in the NCP feed register. "
            "Until now the platform has reported late reports once a week, in the weekly summary of reports received more "
            "than 24 hours after authorisation; that summary continues."]),
        ("The SLA lateness", [
            "Each feed takes an SLA lateness of 1, 2, 4, 6, 8, 12 or 24 hours. The lateness of a report is the time from "
            "the authorisation date and time in the report (OBR-22) to the platform's first acceptance of the report. When "
            "a report's lateness exceeds its feed's SLA lateness, the platform pages the on-call officer of the public "
            "health unit the case is assigned to, with the report and its lateness."]),
        ("Configuration", [
            "Unit administrators enter each feed's SLA lateness in the feed register before go-live. A feed without an SLA "
            "lateness cannot be activated on 4.2. Settings can be changed at any time and take effect from the next report "
            "received."]),
        ("Other changes in 4.2", [
            "Laboratory contacts are shown on the feed register page. The weekly feed volume report adds a column for "
            "rejected messages. Minor fixes to the condition code set 2025 mappings."]),
    ]:
        st.append(Paragraph(head, S["h2"]))
        for p in paras:
            st.append(Paragraph(p, S["body"]))
    _pdf(path, st, "NCP release 4.2 release notes: per-feed SLA alerts", "NCP platform services",
         "NCP 4.2 release notes  |  per-feed SLA alerts  |  for public health units", "2026-10-14")


# ------------------------------------------------------------------------------------------------ feed register (xlsx)
DELIVERY_TEXT = {
    "CRBH": "Real time", "LPG": "Batches every 30 minutes", "HDN": "Batches every hour", "KES": "Real time",
    "HPS": "Batches every 2 hours to 29 February 2024; every 30 minutes from 1 March 2024", "SERL": "Batches every 4 hours",
}


def feed_register(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Feeds"
    rows = []
    for f in P.FEEDS:
        for code, a, b in W.FEED[f]["codes"]:
            rows.append([W.FEED[f]["name"], code, a if a else date(2019, 7, 1), b, "QLD" if f == "KES" else "NSW",
                         "HL7 v2.5 ORU^R01 through the state integration hub", DELIVERY_TEXT[f], None])
    _sheet(ws, ["Feed (laboratory)", "Sending facility", "Active from", "Active to", "State", "Interface", "Delivery",
                "SLA lateness, hours (NCP 4.2)"], rows, widths=[40, 16, 12, 12, 7, 46, 52, 18], wrap=True,
           formats={3: "dd/mm/yyyy", 4: "dd/mm/yyyy"})
    ws2 = wb.create_sheet("Notes")
    for line in ["NCP feed register for the Coast and Ranges district's laboratory feeds, as at 13 November 2026.",
                 "One row per sending facility code, with the dates the code was in use; a laboratory's codes share its "
                 "feed name.",
                 "The SLA lateness column is filled in for the 4.2 release (go-live 1 December 2026).",
                 "Maintained by %s, surveillance data officer." % W.PEOPLE["interface"]]:
        ws2.append([line])
    ws2.column_dimensions["A"].width = 110
    _save(wb, path, W.PEOPLE["interface"], "2021-02-15T00:40:00Z", "2026-11-13T05:12:00Z")


# ------------------------------------------------------------------------------------------------ investigation register
def register(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Register"
    header = ["Investigation", "NCP case", "Condition", "Service", "Room", "Index", "Index person no.", "Onset",
              "Infectious attendance from", "Infectious attendance to", "Notifying laboratory", "Specimen collected",
              "Result authorised", "Notification received", "Exclusion call", "Follow-up to", "Status", "Closed"]
    out = []
    for x in rows:
        out.append([x["investigation"], x["ncp_case"], x["condition"], x["service"], x["room"], x["index_role"],
                    x["index_person"], x["onset"], x["inf_from"], x["inf_to"], x["feed"], x["specimen"],
                    _naive(x["authorised"]), _naive(x["received"]).replace(microsecond=0),
                    _naive(x["call"]).replace(second=0, microsecond=0), x["followup_to"], x["status"],
                    x["closed"]])
    fmt = {k: "dd/mm/yyyy" for k in (8, 9, 10, 12, 16, 18)}
    fmt.update({k: "dd/mm/yyyy hh:mm" for k in (13, 14, 15)})
    _sheet(ws, header, out, widths=[11, 12, 34, 13, 12, 9, 15, 11, 13, 13, 40, 11, 16, 16, 16, 11, 8, 11], formats=fmt)
    ws.auto_filter.ref = "A1:R%d" % (len(out) + 1)
    ws2 = wb.create_sheet("Notes")
    for line in ["Childcare investigation register, %s. Extracted 13 November 2026." % W.UNIT,
                 "Times are New South Wales time. Result authorised is the laboratory's authorisation of the result that "
                 "opened the investigation, as NCP shows it; notification received is the platform's acceptance of that "
                 "report; exclusion call is the duty officer's call to the family and the centre.",
                 "Infectious attendance from and to are the first and last days the index case attended the centre from "
                 "onset to the exclusion.",
                 "Follow-up to is the end of the room's follow-up (GL-CD-11 section 4.4)."]:
        ws2.append([line])
    ws2.column_dimensions["A"].width = 120
    _save(wb, path, W.PEOPLE["manager"], "2021-07-26T22:30:00Z", "2026-11-13T04:05:00Z")


# ------------------------------------------------------------------------------------------------ daily roll (layout B)
def daily_roll(path, data):
    wb = Workbook()
    ws = wb.active
    ws.title = "Roll"
    rows = []
    for inv, recs in data:
        for d, pid, role, a, b in recs:
            rows.append([inv.iid, inv.centre["sid"], inv.room["room"], pid, "Educator" if role == "educator" else "Child",
                         d, _ampm(a), _ampm(b)])
    _sheet(ws, ["Investigation", "Service", "Room", "Attendee", "Type", "Date", "In", "Out"], rows,
           widths=[12, 13, 13, 10, 9, 11, 9, 9], formats={6: "dd/mm/yyyy"})
    ws.auto_filter.ref = "A1:H%d" % (len(rows) + 1)
    _save(wb, path, "Childcare centres (daily roll export)", "2021-08-02T01:15:00Z", "2026-11-04T23:20:00Z")
    return len(rows)


def _ampm(h):
    m = int(round(h * 60))
    hh, mm_ = m // 60, m % 60
    return "%d:%02d %s" % ((hh - 1) % 12 + 1, mm_, "AM" if hh < 12 else "PM")


# ------------------------------------------------------------------------------------------------ dictionary
def dictionary(path, files, fields, notes):
    wb = Workbook()
    ws = wb.active
    ws.title = "Files"
    _sheet(ws, ["file", "contents"], files, widths=[46, 110], wrap=True)
    ws2 = wb.create_sheet("Fields")
    _sheet(ws2, ["file", "field", "description"], fields, widths=[40, 26, 100], wrap=True)
    ws3 = wb.create_sheet("Conventions")
    for line in notes:
        ws3.append([line])
    ws3.column_dimensions["A"].width = 125
    _save(wb, path, W.PEOPLE["interface"], "2023-03-01T02:00:00Z", "2026-11-13T05:30:00Z")
