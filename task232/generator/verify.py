"""Independent verifier: reads only the shipped files under target/ and recomputes the call, its components and the asks.

    python3 task232/generator/verify.py [target_dir]

The governing facts it uses (the incubation ranges, the calling hours, the leaving hour, the tolerance, the allowed
settings, the basis period) are read out of the shipped standard, guideline and release notes and asserted there."""
import csv
import json
import os
import re
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta, timezone
from zoneinfo import ZoneInfo

NSW = ZoneInfo("Australia/Sydney")
UTC = timezone.utc
MONTHS = {m: i + 1 for i, m in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov",
                                          "Dec"])}
FILES = dict(log="ncp_elr_message_log_2022-10_2026-10.csv", extract="ncp_enteric_case_extract_2022-2026.csv",
             feeds="ncp_feed_register.xlsx", register="cc_investigation_register_2021-2026.xlsx",
             followup="cc_contact_followup_2021-2026.csv", att_a="cc_attendance_signin_export.csv",
             att_b="cc_attendance_daily_roll.xlsx", att_c="cc_attendance_portal_export.json",
             bookings="cc_room_bookings.csv", standard="CD-ST-07_laboratory_feed_sla_lateness.docx",
             guideline="GL-CD-11_enteric_ecec_response_guideline.pdf", release="ncp_4-2_release_notes_sla_alerts.pdf")
URGENT = ["Shigellosis", "Shiga toxin-producing E. coli infection", "Cryptosporidiosis"]


# ------------------------------------------------------------------------------------------------ reading
def rows_csv(path):
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def rows_xlsx(path, sheet=None):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = list(next(it))
    out = [dict(zip(head, r)) for r in it if any(v is not None for v in r)]
    wb.close()
    return out


def text_docx(path):
    from docx import Document
    return "\n".join(p.text for p in Document(path).paragraphs)


def text_pdf(path):
    from pypdf import PdfReader
    return "\n".join(pg.extract_text() for pg in PdfReader(path).pages)


def hours(s):
    s = s.strip()
    m = re.fullmatch(r"(\d{1,2}):(\d{2}) ?([AP]M)?", s)
    h, mi, ap = int(m.group(1)), int(m.group(2)), m.group(3)
    if ap:
        h = h % 12 + (12 if ap == "PM" else 0)
    return h + mi / 60.0


def to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    v = str(v)
    m = re.fullmatch(r"(\d{2})-([A-Z][a-z]{2})-(\d{4})", v)
    if m:
        return date(int(m.group(3)), MONTHS[m.group(2)], int(m.group(1)))
    return date.fromisoformat(v[:10])


def hl7_time(s):
    t = datetime.strptime(s[:14], "%Y%m%d%H%M%S")
    sign = 1 if s[14] == "+" else -1
    off = sign * (int(s[15:17]) * 60 + int(s[17:19]))
    return (t - timedelta(minutes=off)).replace(tzinfo=UTC)


def utc_time(s):
    return datetime.strptime(s, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=UTC)


# ------------------------------------------------------------------------------------------------ the governing facts
def facts(target):
    std = text_docx(os.path.join(target, FILES["standard"]))
    gl = re.sub(r"\s+", " ", text_pdf(os.path.join(target, FILES["guideline"])))
    rel = re.sub(r"\s+", " ", text_pdf(os.path.join(target, FILES["release"])))
    assert "07:00 to 21:00 New South Wales time every day" in std
    assert "leaves one hour after the call" in std and "does not attend that day" in std
    assert "do not exceed 1.00" in std
    assert "1 October 2023 to 30 September 2026" in std
    assert "confirmed and probable cases" in std
    assert "1, 2, 4, 6, 8, 12 or 24 hours" in rel and "(1, 2, 4, 6, 8, 12 and 24 hours)" in std
    inc = {}
    for cond, pat in (("Shigellosis", r"Shigellosis 12 hours to 7 days"), ("Shiga toxin-producing E. coli infection",
                      r"STEC infection 1 to 10 days"), ("Cryptosporidiosis", r"Cryptosporidiosis 1 to 12 days")):
        assert re.search(pat, gl), pat
    inc["Shigellosis"] = (0.5, 7)
    inc["Shiga toxin-producing E. coli infection"] = (1, 10)
    inc["Cryptosporidiosis"] = (1, 12)
    return dict(incubation=inc, call_from=7.0, call_to=21.0, leave=1.0, tolerance=1.0,
                settings=[1, 2, 4, 6, 8, 12, 24], basis=(date(2023, 10, 1), date(2026, 9, 30)))


# ------------------------------------------------------------------------------------------------ the investigation files
def load_investigations(target):
    reg = rows_xlsx(os.path.join(target, FILES["register"]), "Register")
    feeds = rows_xlsx(os.path.join(target, FILES["feeds"]), "Feeds")
    att = defaultdict(lambda: defaultdict(dict))       # investigation -> date -> person -> (in, out)
    for r in rows_csv(os.path.join(target, FILES["att_a"])):
        att[r["investigation"]][to_date(r["date"])][r["person_no"]] = (hours(r["sign_in"]), hours(r["sign_out"]))
    for r in rows_xlsx(os.path.join(target, FILES["att_b"]), "Roll"):
        att[r["Investigation"]][to_date(r["Date"])][str(r["Attendee"])] = (hours(r["In"]), hours(r["Out"]))
    with open(os.path.join(target, FILES["att_c"]), encoding="utf-8") as fh:
        js = json.load(fh)
    for ex in js["exports"]:
        for rec in ex["records"]:
            a, b = datetime.fromisoformat(rec["signed_in"]), datetime.fromisoformat(rec["signed_out"])
            att[ex["investigation"]][a.date()][rec["attendee_id"]] = (a.hour + a.minute / 60.0, b.hour + b.minute / 60.0)
    book = defaultdict(dict)
    days = {"Mon": 0, "Tue": 1, "Wed": 2, "Thu": 3, "Fri": 4}
    for r in rows_csv(os.path.join(target, FILES["bookings"])):
        book[r["investigation_id"]][r["person_no"]] = ({days[d] for d in r["booked_days"].split()},
                                                       hours(r["booked_from"]), hours(r["booked_to"]))
    fu = defaultdict(list)
    for r in rows_csv(os.path.join(target, FILES["followup"])):
        fu[r["investigation_id"]].append(r)
    invs = []
    for r in reg:
        x = dict(iid=r["Investigation"], cond=r["Condition"], feed=r["Notifying laboratory"], index=str(r["Index person no."]),
                 onset=to_date(r["Onset"]), inf_from=to_date(r["Infectious attendance from"]),
                 inf_to=to_date(r["Infectious attendance to"]), auth=r["Result authorised"].replace(tzinfo=NSW),
                 closed=r["Status"] == "Closed", att=att[r["Investigation"]], book=book[r["Investigation"]],
                 cases=[c for c in fu[r["Investigation"]] if c["case_status"] in ("Confirmed", "Probable")])
        invs.append(x)
    names = [f["Feed (laboratory)"] for f in feeds]
    feed_names = list(dict.fromkeys(names))
    return invs, feed_names


def others(inv, d):
    return sum(1 for p in inv["att"].get(d, {}) if p != inv["index"])


def child_hours(inv):
    tot = 0.0
    d = inv["inf_from"]
    while d <= inv["inf_to"]:
        day = inv["att"].get(d, {})
        if inv["index"] in day:
            a, b = day[inv["index"]]
            tot += (b - a) * others(inv, d)
        d += timedelta(days=1)
    return tot


def counted(inv, onset, rung, F):
    lo, hi = F["incubation"][inv["cond"]]
    if rung == "all":
        return True
    if rung == "before":
        return inv["onset"] <= onset <= inv["call_date"]
    low = (onset - inv["inf_from"]).days >= lo
    up = (onset - inv["inf_to"]).days <= hi
    return {"lower": low, "upper": up, "window": low and up}[rung]


def rates(invs, rung, F, confirmed_only=False):
    num, den = defaultdict(int), defaultdict(float)
    for inv in invs:
        if not inv["closed"]:
            continue
        den[inv["cond"]] += child_hours(inv)
        for c in inv["cases"]:
            if confirmed_only and c["case_status"] != "Confirmed":
                continue
            if counted(inv, date.fromisoformat(c["onset_date"]), rung, F):
                num[inv["cond"]] += 1
    return {k: num[k] / den[k] for k in den}, dict(num), dict(den)


def call_time(t, F):
    t = t.astimezone(NSW)
    h = t.hour + t.minute / 60.0 + t.second / 3600.0
    if F["call_from"] <= h < F["call_to"]:
        return t
    d = t.date() if h < F["call_from"] else t.date() + timedelta(days=1)
    return datetime(d.year, d.month, d.day, 7, 0, tzinfo=NSW)


def held(call, d, s, e, F):
    cd = call.date()
    if d < cd:
        return e - s
    if d > cd:
        return 0.0
    ch = call.hour + call.minute / 60.0 + call.second / 3600.0
    if ch <= s:
        return 0.0
    return min(ch + F["leave"], e) - s


def added(inv, T, F):
    a = inv["auth"]
    c0 = call_time(a, F)
    cT = call_time((a.astimezone(UTC) + timedelta(hours=T)).astimezone(NSW), F)
    days, s, e = inv["book"][inv["index"]]
    tot = 0.0
    d = c0.date()
    while d <= cT.date():
        if d.weekday() in days:
            h = held(cT, d, s, e, F) - held(c0, d, s, e, F)
            if h > 0:
                tot += h * others(inv, d)
        d += timedelta(days=1)
    return tot


def settings(invs, rate, F, feed_names):
    basis = [i for i in invs if F["basis"][0] <= i["auth"].astimezone(NSW).date() <= F["basis"][1]]
    out = {}
    for f in feed_names:
        rs = [i for i in basis if i["feed"] == f]
        curve = [sum(rate[i["cond"]] * added(i, T, F) for i in rs) / 3.0 for T in F["settings"]]
        ok = [k for k, c in enumerate(curve) if c <= F["tolerance"]]
        k = max(ok) if ok else 0
        out[f] = dict(setting=F["settings"][k], at=curve[k], next=curve[k + 1] if k + 1 < len(curve) else None,
                      curve=curve, n=len(rs))
    return out


def main_call(target, F=None):
    F = F or facts(target)
    invs, feed_names = load_investigations(target)
    reg = {r["Investigation"]: r for r in rows_xlsx(os.path.join(target, FILES["register"]), "Register")}
    for inv in invs:
        inv["call_date"] = reg[inv["iid"]]["Exclusion call"].date()
    out = {}
    for rung in ("all", "lower", "upper", "before", "window"):
        rate, num, den = rates(invs, rung, F)
        out[rung] = dict(rate=rate, num=num, den=den, feeds=settings(invs, rate, F, feed_names))
    return out, invs, feed_names


# ------------------------------------------------------------------------------------------------ the asks
def feed_of_codes(target):
    out = {}
    for f in rows_xlsx(os.path.join(target, FILES["feeds"]), "Feeds"):
        out[f["Sending facility"]] = f["Feed (laboratory)"]
    return out


def load_log(target):
    feed = feed_of_codes(target)
    msgs = rows_csv(os.path.join(target, FILES["log"]))
    for m in msgs:
        m["_feed"] = feed[m["sending_facility"]]
        m["_auth"] = hl7_time(m["authorised_at"])
        m["_recv"] = utc_time(m["received_utc"])
    return msgs


def ask_reports(msgs, key, year=(date(2025, 10, 1), date(2026, 9, 30))):
    """each feed's reports authorised in the year and those first received more than 24 hours after authorisation"""
    first = {}
    for m in sorted(msgs, key=lambda x: (x["_recv"], x["msg_control_id"])):
        if not (year[0] <= m["_auth"].astimezone(NSW).date() <= year[1]):
            continue
        k = key(m)
        if k not in first:
            first[k] = m
    out = defaultdict(lambda: [0, 0])
    for m in first.values():
        out[m["_feed"]][0] += 1
        if (m["_recv"] - m["_auth"]).total_seconds() > 24 * 3600:
            out[m["_feed"]][1] += 1
    return {f: tuple(v) for f, v in out.items()}


REPORT_KEY = lambda m: (m["_feed"], m["patient_ref"], m["test_code"], m["result_status"], m["authorised_at"])


def ask_cases(target, msgs, mode="answer", year=(date(2025, 10, 1), date(2026, 9, 30))):
    """the urgent-condition cases each feed was the first to notify in the year"""
    ext = {r["ncp_case_no"]: r for r in rows_csv(os.path.join(target, FILES["extract"]))}
    urgent = {k for k, r in ext.items() if r["condition"] in URGENT}
    groups = defaultdict(list)
    for m in msgs:
        k = m["ncp_case_no"]
        if k not in urgent:
            continue
        if mode == "answer":
            seen = 0
            while ext[k]["case_status"] == "Merged":
                k = ext[k]["merged_into"]
                seen += 1
                assert seen < 5
        elif mode == "live" and ext[k]["case_status"] == "Merged":
            continue
        groups[k].append(m)
    out = defaultdict(int)
    for k, ms in groups.items():
        if mode == "answer" and ext[k]["case_status"] not in ("Confirmed", "Probable"):
            continue
        f = min(ms, key=lambda x: (x["_recv"], x["msg_control_id"]))
        if year[0] <= f["_recv"].astimezone(NSW).date() <= year[1]:
            out[f["_feed"]] += 1
    return dict(out)


def asks(target):
    msgs = load_log(target)
    return dict(reports=ask_reports(msgs, REPORT_KEY), cases=ask_cases(target, msgs), msgs=msgs)


def main(target):
    out, invs, names = main_call(target)
    g = out["window"]
    print("rates per 1,000 child-hours:", {k: "%.2f" % (1000 * v) for k, v in g["rate"].items()})
    for f in names:
        x = g["feeds"][f]
        print("  %-45s %2d h  %.2f  %s" % (f, x["setting"], x["at"], "%.2f" % x["next"] if x["next"] is not None else "-"))
    a = asks(target)
    for f in names:
        print("  %-45s reports %s late %s cases %s" % (f, a["reports"][f][0], a["reports"][f][1], a["cases"].get(f, 0)))


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "..", "target"))
