"""The independent verifier: reads only the shipped bundle under target/ and the goldens, on its own code path, and
recomputes the call (the shortfall at each record's own turbulence by an Annex M normalisation fitted to Schedule 3,
and by turbulence bins against the fleet), the recovered energy's bin, the default path's answer, and every turbine's
energy lost to downtime and to grid curtailment by OM-07 section 4; then checks the goldens state them."""
import csv
import datetime as dt
import os
import re
import sys
import zoneinfo

import numpy as np
import pandas as pd
import pyarrow.parquet as pq

LON = zoneinfo.ZoneInfo("Europe/London")
T0 = pd.Timestamp("2025-09-30 23:00:00")       # UTC start of the twelve months (local midnight, 1 October 2025)
NI = 52560
SWITCH = pd.Timestamp("2026-02-03 10:00:00")     # site server (UK local) to hosted platform (UTC)
TURB = ["CH%02d" % i for i in range(1, 11)]
AREA = np.pi * 46.5 ** 2
Q = np.linspace(-3.5, 3.5, 57)
QW = np.exp(-0.5 * Q ** 2)
QW /= QW.sum()


def find(target, pat):
    return os.path.join(target, [f for f in os.listdir(target) if re.search(pat, f)][0])


def curve(target):
    from pypdf import PdfReader
    text = " ".join(p.extract_text() for p in PdfReader(find(target, r"schedule_3.*\.pdf$")).pages)
    pts = sorted({(float(a), float(b)) for a, b in re.findall(r"\b(\d{1,2}\.\d)\s+(\d{1,4})\s+0\.\d{3}\b", text)})
    return np.array([a for a, _ in pts]), np.array([b for _, b in pts])


def warranted(cv, v):
    out = np.interp(v, cv[0], cv[1])
    return np.where((v < cv[0][0]) | (v > cv[0][-1]), 0.0, out)


def utc_index(stamps):
    s = pd.Series(pd.to_datetime(stamps)).reset_index(drop=True)
    pre = s <= SWITCH
    loc = s.where(pre)
    second = loc.duplicated(keep="first") & loc.notna()
    a = loc.dt.tz_localize(LON, ambiguous=np.ones(len(s), bool), nonexistent="shift_forward")
    b = loc.dt.tz_localize(LON, ambiguous=np.zeros(len(s), bool), nonexistent="shift_forward")
    conv = a.where(~second, b).dt.tz_convert("UTC").dt.tz_localize(None)
    utc = conv.where(pre, s)
    return (((utc - T0) / pd.Timedelta(minutes=10)).round().astype(int) - 1).to_numpy()


def load(target):
    sc = pq.read_table(find(target, r"^scada_.*\.parquet$")).to_pandas()
    mast = pd.read_csv(find(target, r"^metmast_.*\.csv$"))
    rho = mast["P_76m_Avg"].to_numpy() * 100 / (287.05 * (mast["T_76m_Avg"].to_numpy() + 273.15))
    vm = mast["WS_80m_Avg"].to_numpy() * (rho / 1.225) ** (1 / 3)
    tim = mast["WS_80m_Std"].to_numpy() / np.maximum(mast["WS_80m_Avg"].to_numpy(), 0.5)
    tur = {}
    for t in TURB:
        x = sc[sc.turbine == t].reset_index(drop=True)
        k = utc_index(x.timestamp)
        v = x.wind_speed_mean.to_numpy()
        tur[t] = dict(k=k, v=v * (rho[k] / 1.225) ** (1 / 3), ti=x.wind_speed_sd.to_numpy() / np.maximum(v, 0.5),
                      p=x.active_power_mean.to_numpy(), s=x.state.to_numpy())
    return sc, mast, rho, vm, tim, tur


def lifts_back(target):
    """CH09's present rotor: the lift log's rotor lifted onto its pad, through the layout; in service from the
    work orders' return."""
    import openpyxl
    lay = pd.read_csv(find(target, r"^site_layout.*\.csv$"))
    pad = dict(zip(lay.pad, lay.turbine))
    wb = openpyxl.load_workbook(find(target, r"^crane_lift_log.*\.xlsx$"), read_only=True)
    ups = [r for r in wb.active.iter_rows(min_row=2, values_only=True) if r[0] and "lifted from laydown and fitted" in r[5]
           and "Rotor" in r[5]]
    assert len(ups) == 1 and pad[ups[0][3]] == "CH09"
    wo = pd.read_csv(find(target, r"^work_orders.*\.csv$"), dtype=str, keep_default_na=False)
    r = wo[(wo.turbine == "CH09") & (wo.type == "Major component")].iloc[0]
    back = dt.datetime.strptime(r.returned_to_service, "%Y-%m-%d %H:%M").replace(tzinfo=LON).astimezone(dt.timezone.utc)
    return int(round((pd.Timestamp(back.replace(tzinfo=None)) - T0) / pd.Timedelta(minutes=10)))


def sim_curve(v, ti, cp, ci):
    out = np.empty(len(v))
    for a in range(0, len(v), 100000):
        vv = np.maximum(v[a:a + 100000, None] * (1 + ti[a:a + 100000, None] * Q), 0)
        p = np.minimum(0.5 * 1.225 * AREA * cp * vv ** 3 / 1000.0, 2300.0)
        p = np.where(vv < ci, 0.0, p)
        out[a:a + 100000] = (p * QW).sum(-1)
    return out


def main(target, golden):
    sc, mast, rho, vm, tim, tur = load(target)
    cv = curve(target)
    k9 = lifts_back(target)
    keep = {t: np.isin(tur[t]["s"], (10, 20)) & ((tur[t]["k"] >= k9) if t == "CH09" else True) for t in TURB}
    # the default: method of bins against the warranted curve
    edges = np.arange(-0.25, 30, 0.5)
    cen = edges[:-1] + 0.25
    wc = warranted(cv, cen)
    default = {}
    for t in TURB:
        m = keep[t]
        i = np.clip(np.digitize(tur[t]["v"][m], edges) - 1, 0, len(cen) - 1)
        default[t] = float((wc * np.bincount(i, minlength=len(cen)) - np.bincount(i, tur[t]["p"][m], minlength=len(cen))).sum() / 6000)
    # Annex M: the zero-turbulence curve reproducing Schedule 3 at the mast's median turbulence by bin
    i = np.clip(np.digitize(vm, edges) - 1, 0, len(cen) - 1)
    med = np.array([np.median(tim[i == b]) if (i == b).sum() >= 30 else np.nan for b in range(len(cen))])
    g = np.isfinite(med)
    ref = lambda v: np.interp(v, cen[g], med[g])
    sel = (cv[0] >= 4) & (cv[0] <= 12)
    best = min(((float(((sim_curve(cv[0][sel], ref(cv[0][sel]), cp, ci) - cv[1][sel]) ** 2).sum()), cp, ci)
                for cp in np.arange(0.44, 0.50, 0.001) for ci in (2.5, 2.75, 3.0, 3.25, 3.5)))
    _, cp, ci = best
    matched = {t: float((sim_curve(tur[t]["v"][keep[t]], tur[t]["ti"][keep[t]], cp, ci) - tur[t]["p"][keep[t]]).sum() / 6000)
               for t in TURB}
    call = max(matched, key=matched.get)
    rec = matched[call]
    bin50 = int(np.floor(rec / 50 + 0.5) * 50)
    # turbulence bins against the fleet's median (0.5 m/s by 0.02)
    key = {t: np.floor((tur[t]["v"][keep[t]] + 0.25) / 0.5).astype(int) * 1000 + np.floor(tur[t]["ti"][keep[t]] / 0.02).astype(int)
           for t in TURB}
    means = {t: pd.Series(tur[t]["p"][keep[t]]).groupby(key[t]).agg(["mean", "size"]) for t in TURB}
    binned = {}
    for t in TURB:
        others = pd.concat([means[o]["mean"].where(means[o]["size"] >= 3).rename(o) for o in TURB if o != t], axis=1)
        n = pd.concat([means[o]["size"].rename(o) for o in TURB if o != t], axis=1).fillna(0).sum(1)
        refm = others.median(axis=1)
        r = refm.reindex(key[t]).to_numpy()
        nn = n.reindex(key[t]).fillna(0).to_numpy()
        okc = (nn >= 10) & np.isfinite(r)
        binned[t] = float((r - tur[t]["p"][keep[t]])[okc].sum() / 6000 / okc.mean())
    # lost production, OM-07 section 4
    pot = warranted(cv, vm)
    anm = pd.read_csv(find(target, r"^anm_.*\.csv$"))
    force = np.zeros(NI, bool)
    for a, b in zip(anm.start_utc, anm.end_utc):
        ka = int(round((pd.Timestamp(a[:19]) - T0) / pd.Timedelta(minutes=10)))
        kb = int(round((pd.Timestamp(b[:19]) - T0) / pd.Timedelta(minutes=10)))
        force[ka:kb] = True
    wo = pd.read_csv(find(target, r"^work_orders.*\.csv$"), dtype=str, keep_default_na=False)

    def when(day, s):
        if not s:
            return None
        d = dt.date.fromisoformat(s[:10]) if len(s) > 5 else dt.date.fromisoformat(day)
        hh, mm = s[-5:].split(":")
        u = dt.datetime(d.year, d.month, d.day, int(hh), int(mm), tzinfo=LON).astimezone(dt.timezone.utc)
        return int(round((pd.Timestamp(u.replace(tzinfo=None)) - T0) / pd.Timedelta(minutes=10)))
    D, C = {}, {}
    for t in TURB:
        st = np.full(NI, -1)
        p = np.zeros(NI)
        st[tur[t]["k"]] = tur[t]["s"]
        p[tur[t]["k"]] = tur[t]["p"]
        silent = st < 0
        down = np.isin(st, (30, 40, 41))
        for _, r in wo[wo.turbine == t].iterrows():
            a, b = ((when(r.date, r.isolated_from), when(r.date, r.isolated_to)) if r.type == "Annual service"
                    else (when(r.date, r.turbine_stopped), when(r.date, r.returned_to_service)))
            down[a:b] |= silent[a:b]
        # a silent stretch with the grid state on both sides is the grid loss
        neither = np.isin(st, (21, 22, 61))
        runs = np.flatnonzero(np.diff(np.concatenate([[0], silent.astype(int), [0]])))
        for a, b in zip(runs[::2], runs[1::2]):
            if a > 0 and b < NI and st[a - 1] == 61 and st[b] == 61:
                neither[a:b] = True
        lp = np.maximum(pot - p, 0) / 6000
        D[t] = float(lp[down].sum())
        C[t] = float(lp[force & ~down & ~neither].sum())
    # the goldens
    import openpyxl
    from pypdf import PdfReader
    wb = openpyxl.load_workbook(os.path.join(golden, "lost_production_2025_26.xlsx"), read_only=True)
    rows = {r[0]: r for r in wb["Lost production 2025-26"].iter_rows(min_row=2, values_only=True) if r[0] in TURB}
    text = re.sub(r"\s+", " ", " ".join(p.extract_text() for p in PdfReader(os.path.join(golden, "le_repair_campaign_2027.pdf")).pages))
    errs = []
    if call != "CH04" or max(binned, key=binned.get) != "CH04":
        errs.append("call %s / %s" % (call, max(binned, key=binned.get)))
    if bin50 != 100 or int(np.floor(binned["CH04"] / 50 + 0.5) * 50) != 100:
        errs.append("recovered energy bin %d" % bin50)
    if max(default, key=default.get) != "CH07":
        errs.append("default %s" % max(default, key=default.get))
    if "campaign on %s" % call not in text or "recover %d MWh to the nearest 50 MWh" % bin50 not in text:
        errs.append("paper does not state the call and its bin")
    for t in TURB:
        if rows[t][1] != int(round(D[t])) or rows[t][2] != int(round(C[t])):
            errs.append("%s: workbook %s/%s, verifier %.1f/%.1f" % (t, rows[t][1], rows[t][2], D[t], C[t]))
    if errs:
        raise SystemExit("verify FAILED: " + "; ".join(errs))
    print("verify: call %s (%.1f MWh at matched turbulence, bins %.1f; %d to the nearest 50), default %s, runner-up %s; "
          "downtime %d and curtailment %d MWh over the farm, every turbine's figures as the workbook states" % (
              call, rec, binned["CH04"], bin50, max(default, key=default.get),
              sorted(matched, key=lambda x: -matched[x])[1], round(sum(D.values())), round(sum(C.values()))))


if __name__ == "__main__":
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "target"),
         sys.argv[2] if len(sys.argv) > 2 else os.path.join(here, "golden"))
