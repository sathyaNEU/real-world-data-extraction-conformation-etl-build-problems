"""Write the shipped files from the clean world, each in its producer's own layout, with the mess its system makes:
the van tablet's UTC clock, its mid-year change of hostel site names and its double syncs, the register's two systems
across the October 2024 go-live, the laboratory's change of RNA assay and re-sent results, the two hostel providers'
lists, and the pharmacy system's 'New' rows for returning clients."""
import csv
import datetime as dt
import hashlib
import os

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import params as P
import postprocess
import world as Wm

D = dt.date
DAY = dt.timedelta(days=1)


def _w(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\r\n")
        w.writerow(header)
        w.writerows(rows)


def _h(*parts, n=6):
    return hashlib.sha256("|".join(map(str, parts)).encode()).hexdigest().upper()[:n]


def _local(ts):
    return ts.replace(tzinfo=dt.timezone.utc).astimezone(Wm.UK)


# ---------------------------------------------------------------------------- the van tablet
WORKERS = ["JM", "SK", "TD", "AH", "RB"]


def van(W, path, rng):
    rows = []
    for c in W.contacts:
        loc = _local(c["when"])
        setting = P.TABLET_OLD[c["setting"]]
        site = "Recovery hub, Gorse Street" if c["setting"] == "DS" else c["site"]
        if c["setting"] == "HH" and loc.date() >= P.TABLET_UPDATE:
            site = P.HOSTEL_SHORT[site]
        row = ["BTV-" + _h(c["pid"], c["when"]), c["when"].strftime("%Y-%m-%d %H:%M:%S"),
               "TAB-0%d" % (2 if c["setting"] in ("HH", "AP") else 5), setting, site, c["pid"],
               "Accepted" if c["accepted"] else "Declined", c["ab"], c["rna"], WORKERS[int(rng.integers(len(WORKERS)))]]
        rows.append(row)
        if rng.random() < 0.008:
            rows.append(list(row))              # the tablet synced the same record twice
    _w(path, ["contact_ref", "recorded_at_utc", "tablet", "setting", "site", "client_pseudonym", "test_outcome",
              "antibody_result", "rna_result", "peer_worker"], rows)


# ---------------------------------------------------------------------------- register and dispensing
REGIMEN = {8: "Glecaprevir/pibrentasvir 8 weeks", 12: "Sofosbuvir/velpatasvir 12 weeks"}
DISPENSARY = {"HEP": "Royal Infirmary outpatient pharmacy", "SHC": "Royal Infirmary outpatient pharmacy",
              "COMM": "Kestle Road clinic pharmacy", "FRC": "Recovery hub dispensary", "HHC": "Riverside Pharmacy"}
VAN_SOURCE = ["Brannock Trust van", "Brannock van", "BT outreach van"]


def course_refs(W, rng):
    """old-system references before the go-live, new-system ones after; courses open at the go-live were closed on the
    old system and re-opened on the new one, pointing back to the old reference"""
    used, seq = set(), {}
    for c in sorted(W.courses, key=lambda c: (c["start"], c["pid"])):
        if c["start"] < P.EPR_GO_LIVE:
            y = c["start"].year
            k = int(rng.integers(1, 9999))
            while (y, k) in used:
                k = int(rng.integers(1, 9999))
            used.add((y, k))
            c["ref"] = "FHN/%d/%04d" % (y, k)
            if c["planned_end"] >= P.EPR_GO_LIVE:
                c["new_ref"] = "C%07d" % int(rng.integers(2400000, 2499999))
        else:
            k = "C%07d" % int(rng.integers(2400000, 2699999))
            while k in seq:
                k = "C%07d" % int(rng.integers(2400000, 2699999))
            seq[k] = 1
            c["ref"] = k


def register(W, path, rng):
    rows = []
    for c in sorted(W.courses, key=lambda c: (c["start"], c["ref"])):
        # the end date is entered at the end-of-treatment review, from the prescription's planned length
        end = c["planned_end"].isoformat() if c["planned_end"] <= P.AS_OF else ""
        new = c["start"] >= P.EPR_GO_LIVE
        code = c["clinic"] if (new or c["clinic"] == "PHARM") else P.CLINIC_OLD[c["clinic"]]
        role = {"Consultant": "Consultant", "Nurse": "Clinical nurse specialist", "Pharmacist":
                "Pharmacist independent prescriber"}[c["role"]]
        src = VAN_SOURCE[int(rng.integers(3))] if c["source"] == "Brannock Trust van" else c["source"]
        if "new_ref" in c:
            rows.append([c["ref"], c["pid"], code, role, c["site"], src, REGIMEN[c["weeks"]], c["start"].isoformat(),
                         (P.EPR_GO_LIVE - DAY).isoformat(), ""])
            rows.append([c["new_ref"], c["pid"], c["clinic"], role, c["site"], src, REGIMEN[c["weeks"]],
                         P.EPR_GO_LIVE.isoformat(), end, c["ref"]])
        else:
            rows.append([c["ref"], c["pid"], code, role, c["site"], src, REGIMEN[c["weeks"]], c["start"].isoformat(),
                         end, ""])
    rows.sort(key=lambda r: (r[7], r[0]))
    _w(path, ["course_ref", "pseudonym", "clinic_code", "prescriber_role", "prescribing_site", "referral_source",
              "regimen", "start_date", "end_date", "migrated_from"], rows)


def dispensing(W, path, rng, names):
    rows, used = [], set()
    for c in W.courses:
        for d in c["packs"]:
            if d < D(2024, 4, 1) or d > P.AS_OF:
                continue
            ref = c["ref"] if ("new_ref" not in c or d < P.EPR_GO_LIVE) else c["new_ref"]
            k = "DP%08d" % int(rng.integers(10 ** 7, 10 ** 8))
            while k in used:
                k = "DP%08d" % int(rng.integers(10 ** 7, 10 ** 8))
            used.add(k)
            ph = names[c["site"]] if c["clinic"] == "PHARM" else DISPENSARY[c["clinic"]]
            prod = REGIMEN[c["weeks"]].rsplit(" ", 2)[0]
            rows.append([k, ref, c["pid"], d.isoformat(), prod, "84 tablets" if c["weeks"] == 8 else "28 tablets", 28, ph])
    rows.sort(key=lambda r: (r[3], r[0]))
    _w(path, ["dispensing_ref", "course_ref", "pseudonym", "dispensed_on", "product", "quantity", "days_supply",
              "dispensed_by"], rows)


# ---------------------------------------------------------------------------- the laboratory
GP = ["Ashgrove Health Centre", "Halden Medical Practice", "Quay Street Surgery", "Okeby Road Surgery",
      "Whitcliffe Green Practice", "Stanmoor Family Practice", "Northgate Medical Centre", "Carrbrook Surgery",
      "Highfield Avenue Surgery", "Mill Road Surgery", "Vale Street Medical Centre", "Moorside Practice",
      "Lowmoor Health Centre", "Pennant Way Surgery", "Calder Street Surgery", "Coldharbour Medical Practice"]
LOC = {"HEP": "Royal Infirmary Hepatology", "COMM": "Kestle Road Community Clinic", "FRC": "Fenwold Recovery",
       "HHC": "Riverside Homeless Health", "SHC": "Sexual Health Fenwold", "Drug service": "Fenwold Recovery",
       "Sexual health": "Sexual Health Fenwold", "Prison": "HMP Fenwold Healthcare",
       "Antenatal": "Royal Infirmary Antenatal", "Emergency department": "Royal Infirmary ED",
       "PHARM": "Community Pharmacy HCV Service"}


def lab(W, path, rng):
    rows, used = [], set()
    for x in W.lab:
        if x["date"] > P.AS_OF:
            continue
        k = "F%s%07d" % (x["date"].strftime("%y"), int(rng.integers(10 ** 6, 10 ** 7)))
        while k in used:
            k = "F%s%07d" % (x["date"].strftime("%y"), int(rng.integers(10 ** 6, 10 ** 7)))
        used.add(k)
        if x["kind"] == "AB":
            code, name = "HCVAB", "Hepatitis C antibody"
            res, vl = ("Reactive" if x["result"] == "Reactive" else "Not reactive"), ""
        else:
            new = x["date"] >= P.ASSAY_CHANGE
            code, name = ("HCVQNT", "HCV RNA quantitative") if new else ("HCVPCR", "HCV RNA (PCR)")
            res = {"Detected": "Detected", "Not detected": "Not detected", "Detected <12": "Detected <12 IU/mL"}[x["result"]]
            vl = x["vl"] if x["result"] == "Detected" else ""
        loc = GP[int(rng.integers(len(GP)))] if x["loc"] == "GP" else LOC[x["loc"]]
        auth = dt.datetime.combine(x["date"] + int(rng.integers(1, 5)) * DAY, dt.time(0)) + dt.timedelta(
            seconds=int(rng.integers(7 * 3600, 19 * 3600)))
        row = [k, x["pid"], x["date"].isoformat(), code, name, res, vl, loc, auth.strftime("%Y-%m-%dT%H:%M:%SZ")]
        rows.append(row)
        if rng.random() < 0.012:
            rows.append(list(row))              # the result re-sent to the network's feed
    _w(path, ["lab_number", "patient_pseudonym", "sample_collected", "test_code", "test_name", "result",
              "viral_load_iu_ml", "requesting_location", "authorised_utc"], rows)


# ---------------------------------------------------------------------------- settings' records
def census(W, path, rng):
    wb = Workbook()
    a = wb.active
    a.title = "FHT residents"
    a.append(["Fenwold Housing Trust: residents at 30 September 2026"])
    a.append([])
    a.append(["Resident ref", "Scheme", "Room", "Moved in", "Support level"])
    cs = wb.create_sheet("CSH 30.09.26")
    cs.append(["Hostel", "Name ref", "Date moved in", "Key worker"])
    res = [r for r in W.hostel_res if r["move_out"] is None]
    res.sort(key=lambda r: (r["hostel"], r["move_in"], r["pid"]))
    for r in res:
        if r["provider"] == "Fenwold Housing Trust":
            room = "%s%02d" % ("GF12"[int(rng.integers(4))], int(rng.integers(1, 30)))
            a.append([r["pid"], r["hostel"], room, r["move_in"], ["Low", "Medium", "High"][int(rng.choice(3, p=[.3, .5, .2]))]])
            a.cell(a.max_row, 4).number_format = "dd/mm/yyyy"
        else:
            cs.append([r["hostel"], r["pid"], r["move_in"].strftime("%d/%m/%Y"),
                       ["KM", "PO", "DW", "LF", "SN", "RT"][int(rng.integers(6))]])
    a["A1"].font = Font(bold=True)
    for c in a[3]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E1F2")
    for c in cs[1]:
        c.font = Font(bold=True)
    for col, wd in zip("ABCDE", (16, 20, 8, 12, 14)):
        a.column_dimensions[col].width = wd
    for col, wd in zip("ABCD", (24, 14, 15, 11)):
        cs.column_dimensions[col].width = wd
    wb.save(path)
    postprocess.scrub(path, "Fenwold Housing Trust", "2026-10-02")
    postprocess.finish_office(path, "2026-10-01T10:12:00Z", "2026-10-02T08:40:00Z", "Fenwold Housing Trust")


def _occupancy(stays, m0, m1, unit):
    rows = []
    for y, m in Wm.months(m0, m1):
        a, b = Wm.month_days(y, m)
        for name in unit:
            here = [s for s in stays if s[0] == name]
            res = sum(1 for s in here if s[1] <= b and (s[2] is None or s[2] > b))
            arr = sum(1 for s in here if a <= s[1] <= b)
            dep = sum(1 for s in here if s[2] is not None and a <= s[2] <= b)
            rows.append(((y, m), name, res, arr, dep))
    return rows


def hostel_returns(W, path):
    prov = {h[0]: (h[1], h[2]) for h in P.HOSTELS}
    stays = [(r["hostel"], r["move_in"], r["move_out"]) for r in W.hostel_res]
    rows = [["%d-%02d" % ym, prov[n][0], n, prov[n][1], res, arr, dep]
            for ym, n, res, arr, dep in _occupancy(stays, (2023, 10), (2026, 9), [h[0] for h in P.HOSTELS])]
    _w(path, ["month", "provider", "hostel", "bed_spaces", "residents_at_month_end", "arrivals", "departures"], rows)


def ap_returns(W, path):
    beds = dict(P.AP_PREMISES)
    stays = [(r["premises"], r["arrive"], r["leave"]) for r in W.ap_res]
    rows = [[D(*ym, 1).strftime("%b-%y"), n, beds[n], res, arr, dep]
            for ym, n, res, arr, dep in _occupancy(stays, (2023, 10), (2026, 9), list(beds))]
    _w(path, ["Month", "Approved premises", "Beds", "Residents at month end", "Arrivals", "Departures"], rows)


def ds_episodes(W, path, rng):
    rows = []
    for (y, m), n in sorted(W.ds_monthly.items()):
        k = rng.multinomial(n, [0.47, 0.23, 0.09, 0.21])
        rows.append(["%d-%02d" % (y, m), n] + [int(v) for v in k])
    _w(path, ["month", "new_treatment_episodes", "opiate", "non_opiate", "non_opiate_and_alcohol", "alcohol_only"], rows)


def nsp_registrations(W, path, rng):
    rows, used = [], set()
    for r in sorted(W.nsp_reg, key=lambda r: (r["date"], r["client"])):
        k = "REG%07d" % int(rng.integers(10 ** 6, 10 ** 7))
        while k in used:
            k = "REG%07d" % int(rng.integers(10 ** 6, 10 ** 7))
        used.add(k)
        rows.append([k, r["client"], r["pharmacy"], r["date"].isoformat()])
    _w(path, ["registration_id", "client_code", "pharmacy_code", "registered_on"], rows)


HOURS = ["Mon-Fri 09:00-18:00, Sat 09:00-13:00", "Mon-Sat 08:30-18:30", "Mon-Fri 09:00-17:30",
         "Mon-Sun 08:00-22:00", "Mon-Fri 08:30-18:00, Sat 09:00-17:00", "Mon-Sat 07:00-23:00, Sun 10:00-16:00"]


def pharmacies(W, path, rng):
    rows = [[p["code"], p["name"], p["address"], p["district"], p["level"], HOURS[int(rng.integers(len(HOURS)))]]
            for p in sorted(W.pharmacies, key=lambda p: p["name"])]
    _w(path, ["ods_code", "pharmacy", "address", "postcode_district", "nsp_service", "opening_hours"], rows)


ADDR = {"HEP": "Outpatients C, Fenwold Royal Infirmary, Infirmary Road", "COMM": "Kestle Road Health Centre, 77 Kestle Road",
        "FRC": "The Recovery Hub, 12 Gorse Street", "HHC": "Riverside Health Centre, Quay Street",
        "SHC": "Fenwold Sexual Health, Princess Wing, Fenwold Royal Infirmary"}


def clinics(path):
    rows = [[c, P.CLINIC_OLD[c], P.CLINIC_NAME[c][0].upper() + P.CLINIC_NAME[c][1:], ADDR[c], DISPENSARY[c]]
            for c in P.CLINICS]
    _w(path, ["clinic_code", "code_before_1_oct_2024", "clinic", "address", "dispensing_pharmacy"], rows)


# ---------------------------------------------------------------------------- two files the solution never reads
def fibroscan(W, path, rng):
    wb = Workbook()
    ws = wb.active
    ws.title = "Scans 2026"
    ws.append(["Scan date", "Setting", "Site", "Client pseudonym", "Median kPa", "IQR/median %", "Valid readings",
               "Probe", "Operator"])
    found = [c for c in W.contacts if c["rna"] == "Detected"] + [c for c in W.contacts if c["ab"] == "Reactive"
                                                                  and c["rna"] != "Detected"][:40]
    for c in sorted(found, key=lambda c: c["when"]):
        if rng.random() < 0.18:
            continue
        kpa = round(float(rng.lognormal(1.95, 0.55)), 1)
        ws.append([_local(c["when"]).date(), P.SHORT[c["setting"]], "Recovery hub" if c["setting"] == "DS" else c["site"],
                   c["pid"], kpa, int(rng.integers(6, 38)), int(rng.integers(8, 11)), "M" if rng.random() < 0.8 else "XL",
                   ["AH", "RB"][int(rng.integers(2))]])
        ws.cell(ws.max_row, 1).number_format = "dd/mm/yyyy"
    for c in ws[1]:
        c.font = Font(bold=True)
    wb.save(path)
    postprocess.scrub(path, P.ORG, "2026-10-05")
    postprocess.finish_office(path, "2026-02-09T11:20:00Z", "2026-10-05T16:02:00Z", P.ORG)


def costs(W, path, rng):
    wb = Workbook()
    ws = wb.active
    ws.title = "Running costs"
    ws.append(["Month", "Setting", "Van days", "Nurse sessions", "Peer worker hours", "Fuel (litres)",
               "POC cartridges used", "Consumables (GBP)", "Staff cost (GBP)"])
    days = Wm.rota_days()
    for m in range(1, 10):
        for s in ("HH", "NSP", "DS", "AP"):
            n = sum(1 for d in days[s] if d.month == m)
            tests = sum(1 for c in W.contacts if c["setting"] == s and c["accepted"] and _local(c["when"]).month == m)
            ws.append([D(2026, m, 1).strftime("%b %Y"), P.SHORT[s], n, n, round(n * float(rng.uniform(6.5, 7.5)), 1),
                       round(n * float(rng.uniform(14, 31)), 1), tests + int(rng.integers(0, 4)),
                       round(tests * 6.4 + n * float(rng.uniform(18, 26)), 2), round(n * float(rng.uniform(612, 655)), 2)])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True)
    wb.save(path)
    postprocess.scrub(path, P.ORG, "2026-10-05")
    postprocess.finish_office(path, "2026-01-30T09:00:00Z", "2026-10-05T15:47:00Z", P.ORG)


def write_all(W, target):
    import docs
    rng = np.random.default_rng(P.SEED + 99)
    f = lambda k: os.path.join(target, P.FILES[k])
    names = {p["code"]: p["name"] for p in W.pharmacies}
    course_refs(W, rng)
    van(W, f("van"), rng)
    register(W, f("register"), rng)
    dispensing(W, f("dispensing"), rng, names)
    lab(W, f("lab"), rng)
    census(W, f("census"), rng)
    hostel_returns(W, f("hostel_returns"))
    ap_returns(W, f("ap_returns"))
    ds_episodes(W, f("ds_episodes"), rng)
    nsp_registrations(W, f("nsp_reg"), rng)
    pharmacies(W, f("pharmacies"), rng)
    clinics(f("clinics"))
    fibroscan(W, f("fibroscan"), rng)
    costs(W, f("costs"), rng)
    docs.plan(f("plan"))
    docs.pathway(f("pathway"))
    docs.intentions(f("intentions"))
