"""task304 documents: the 2027 plan paper, the 2024 pilot report, the arrears and possessions policy and the board's
risk appetite statement (pdf), the product transfer rate card (xlsx)."""
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.4, leading=13.0, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=13.2, leading=17, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.2, leading=14,
                        spaceBefore=8, spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8.0, leading=10.6, textColor=colors.HexColor("#444444"))
    return body, h1, h2, small


def _footer(text):
    def foot(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.4)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(18 * mm, 10 * mm, text)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, "%d" % d.page)
        canvas.restoreState()
    return foot


def _table(rows, widths, size=8.4):
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=size, leading=size + 2.2)
    head = ParagraphStyle("head", parent=cell, fontName="Helvetica-Bold")
    rows = [[Paragraph(str(v), head if r == 0 else cell) for v in row] for r, row in enumerate(rows)]
    t = Table(rows, colWidths=[w * mm for w in widths], repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9A9A9A")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E3E8EC")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


# the times each paper was first saved and last saved on its date, and the UK offset that day
SAVED = {"2026-10-09": ("101412", "152703", "+01'00'"), "2026-07-14": ("094155", "160238", "+01'00'"),
         "2026-03-20": ("112007", "115844", "+00'00'"), "2026-02-24": ("140531", "171209", "+00'00'")}


def _pdf(path, title, subject, story, foot, stamp):
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title=title, author=P.ORG, creator=P.ORG, subject=subject,
                            invariant=1)
    doc.build(story, onFirstPage=_footer(foot), onLaterPages=_footer(foot))
    postprocess.scrub(path, P.ORG, stamp)
    postprocess.pdf_clock(path, *SAVED[stamp])


def plan(path, rates):
    body, h1, h2, small = _styles()
    s = [Paragraph("2027 plan: placing the Financial Support Team", h1),
         Paragraph("Paper to the Credit Committee from %s, Head of Collections and Recoveries. 9 October 2026. "
                   "Reference CC/2026/41." % P.PEOPLE["head"], small),
         Paragraph("1. Decision", h2),
         Paragraph("The 2027 plan funds a team of eight financial support officers for the whole of 2027, starting on "
                   "4 January, within the collections budget the Board approved in September: \u00a3436,000 for the "
                   "year, salaries, training and the case system licence included. The team will work one of our four "
                   "lending books: First-time buyer 2019-21, High LTV 2022-24, Shared ownership or Buy to let. The "
                   "Committee is asked to choose the book; it will see the team's first results with the half-year "
                   "arrears report in July 2027.", body),
         Paragraph("2. Who the team works with", h2),
         Paragraph("In the chosen book the team takes on each borrower whose account reaches two monthly payments in "
                   "arrears during 2027, counting from an account with fewer than two, in the first month of the "
                   "year it does so, as in the pilot. The book is the accounts on it at 30 September 2026, each "
                   "counted for the whole of 2027 with no allowance for accounts that close during the year; accounts "
                   "already in litigation or with a receiver appointed at that date are left with the recoveries "
                   "team. Payments in arrears are measured as in section 2 of the Arrears and Possessions Policy. The "
                   "team's capacity, about 1,900 new cases a year, is more than any one book will produce, so no "
                   "borrower who qualifies is turned away.", body),
         Paragraph("3. How the team is valued", h2),
         Paragraph("The team is valued by the possessions it avoids among the borrowers it takes on in 2027. For "
                   "each book we plan on the share of borrowers reaching two payments in arrears who go on to "
                   "possession within 18 months, taken from the book's own experience of the borrowers who first "
                   "reached two payments in arrears from January 2022 to March 2025 (the latest whose 18 months had "
                   "run by September 2026). The pilot showed the team cuts that chance by %d per cent for the borrowers "
                   "it works with, so the possessions avoided in a book are %d per cent of its expected 2027 "
                   "borrowers taken on times "
                   "its planning rate. The costs are the same whichever book is chosen, so the book with the most "
                   "possessions avoided is the one to choose." % (round(100 * P.EFFECT), round(100 * P.EFFECT)), body),
         _table([["Book", "Planning rate: possession within 18 months of reaching two payments in arrears"]] +
                [[b, "%.1f per cent" % (100 * rates[b])] for b in P.ORDER], [52, 118]),
         Spacer(1, 4),
         Paragraph("4. Records", h2),
         Paragraph("The mortgage book at 30 September 2026, the monthly payment history since January 2019, the "
                   "possession register and the litigation actions reported by our panel firms are the records of "
                   "record. The 2027 product transfer rates are those on the rate card of 1 October 2026.", body)]
    _pdf(path, "2027 plan: placing the Financial Support Team", "Credit Committee paper CC/2026/41", s,
         "%s  |  Credit Committee  |  CC/2026/41" % P.ORG, "2026-10-09")


def pilot(path, res):
    body, h1, h2, small = _styles()
    t = res
    cut4 = int(round((1 - (t["team_four"] / t["team_n"]) / (t["usual_four"] / t["usual_n"])) * 20)) * 5
    rt, ru = t["team_claim"] / t["team_n"], t["usual_claim"] / t["usual_n"]
    claims = ("fewer of them had a claim issued" if rt < ru else "about as many of them had a claim issued"
              if abs(rt - ru) < 0.005 else "slightly more of them had a claim issued")
    s = [Paragraph("Financial Support Team pilot 2024: final report", h1),
         Paragraph("Prepared by %s, Financial Support Team Manager, for the Arrears Committee. 14 July 2026."
                   % P.PEOPLE["fst"], small),
         Paragraph("1. What we tested", h2),
         Paragraph("From January to December 2024 a team of four financial support officers worked with borrowers "
                   "alongside the usual collections process. A borrower entered the pilot in the month their account "
                   "first reached two monthly payments in arrears, counting from an account with fewer than two "
                   "(accounts with a claim for possession or an LPA receiver appointment already on record were not "
                   "eligible). Eligible "
                   "borrowers in all four lending books were allocated on the last digit of the account number: even "
                   "to the support team, odd to usual collections. The officers make a full income and expenditure "
                   "assessment, refer to free debt advice, arrange benefit checks and set up the support options in "
                   "the Arrears and Possessions Policy where they fit.", body),
         Paragraph("2. What we measured", h2),
         Paragraph("The primary outcome is an account reaching four monthly payments in arrears within 12 months of "
                   "the month it reached two, the last point before a claim for possession can be issued (Arrears and "
                   "Possessions Policy, section 5); the secondary outcome is a claim for possession issued within the "
                   "same 12 months. Follow-up ran to the end of December 2025 for every pilot borrower.", body),
         _table([["Group", "Borrowers", "Four payments in arrears within 12 months", "Claim issued within 12 months"],
                 ["Support team (even)", "%d" % t["team_n"], "%d (%.1f%%)" % (t["team_four"], 100 * t["team_four"] / t["team_n"]),
                  "%d (%.1f%%)" % (t["team_claim"], 100 * t["team_claim"] / t["team_n"])],
                 ["Usual collections (odd)", "%d" % t["usual_n"], "%d (%.1f%%)" % (t["usual_four"], 100 * t["usual_four"] / t["usual_n"]),
                  "%d (%.1f%%)" % (t["usual_claim"], 100 * t["usual_claim"] / t["usual_n"])]], [46, 26, 50, 50]),
         Paragraph("3. Result", h2),
         Paragraph("Borrowers allocated to the team reached four payments in arrears about %d per cent less often "
                   "than those left with usual collections, and %s, though claims are too few to measure that effect "
                   "on their own. For planning we take a %d per cent cut in the chance of possession, under half the cut in "
                   "arrears, for every borrower the team works with, until the possessions themselves can be counted."
                   % (cut4, claims, round(100 * P.EFFECT)), body),
         Paragraph("4. Capacity", h2),
         Paragraph("Each officer carried about 60 open cases at a time and cases stayed open for a median of three "
                   "months, so each officer can take on about 240 new cases a year.", body)]
    _pdf(path, "Financial Support Team pilot 2024: final report", "Arrears Committee", s,
         "%s  |  Financial Support Team  |  pilot report" % P.ORG, "2026-07-14")


def policy(path):
    body, h1, h2, small = _styles()
    s = [Paragraph("Arrears and Possessions Policy", h1),
         Paragraph("Collections and Recoveries. Version 7, effective 1 April 2026. Owner: %s, Head of Collections "
                   "and Recoveries." % P.PEOPLE["head"], small),
         Paragraph("1. Scope", h2),
         Paragraph("All residential and buy-to-let mortgages held by the Society.", body),
         Paragraph("2. Measuring arrears", h2),
         Paragraph("Arrears are the contractual payments due and not received. Payments in arrears is the arrears "
                   "balance divided by the current contractual monthly payment, rounded down to a whole number of "
                   "payments. Reporting uses the position at the month end.", body),
         Paragraph("3. Contact", h2),
         Paragraph("We contact a borrower within five working days of a missed payment, by letter, telephone and "
                   "text. Borrowers who tell us they are in difficulty before they miss a payment are referred to the "
                   "same team.", body),
         Paragraph("4. Support options", h2),
         Paragraph("In line with MCOB 13.3.4A we consider every option that may help the borrower: an arrangement to "
                   "pay the arrears over time; a temporary switch to interest only (up to six months); a reduced "
                   "payment for a set period; a term extension where the borrower's age and the term limits allow; a "
                   "product transfer to a lower rate; capitalisation of the arrears (once, after six months of keeping "
                   "to an arrangement); and, in exceptional cases, deferral of interest. Since July 2023 a borrower who "
                   "is up to date may extend their term or switch to interest only for six months without a fresh "
                   "affordability assessment, as set out in the Mortgage Charter.", body),
         Paragraph("5. Litigation", h2),
         Paragraph("We do not issue a claim for possession until at least five payments are in arrears, the "
                   "pre-action protocol has been followed and no arrangement has been kept. Claims are conducted by "
                   "our two panel firms in the county court hearing centre for the property. The court may make an "
                   "outright order or an order suspended on terms the borrower keeps to, may vary an order at a later "
                   "hearing, and may set an order aside. A claim stays open until it is withdrawn, the order is "
                   "discharged or set aside, or the property is taken into possession. A possession order stays in "
                   "force until it is varied (the new order takes its place), discharged, set aside or executed by "
                   "eviction. Where the borrower "
                   "breaks the terms or an outright order expires we may apply for a warrant; the County Court "
                   "bailiff carries out the eviction and the Society takes possession that day.", body),
         Paragraph("6. Buy to let", h2),
         Paragraph("Where a buy-to-let mortgage is five or more payments in arrears we may appoint a receiver under "
                   "the Law of Property Act 1925 to collect rent or sell the property. The receiver takes control of "
                   "the property without a court order; the register records that day as the possession. Receivers "
                   "sell and remit the proceeds under their own reference.", body),
         Paragraph("7. Shared ownership", h2),
         Paragraph("We notify the housing association landlord when a shared ownership mortgage reaches two payments "
                   "in arrears. The landlord may buy back the borrower's share or nominate a buyer before "
                   "possession.", body),
         Paragraph("8. Selling a property in possession", h2),
         Paragraph("A property taken into possession is sold on the open market and the account stays open until "
                   "the sale completes. For reporting, the loss on a possession is the debt at possession (the "
                   "balance, the arrears and the legal costs charged to the account) less the sale price net of the "
                   "agent's and conveyancing fees; where the sale raises more, the surplus is paid to the borrower and "
                   "the loss is nil.", body)]
    _pdf(path, "Arrears and Possessions Policy", "Collections and Recoveries policy, version 7", s,
         "%s  |  Arrears and Possessions Policy  |  v7" % P.ORG, "2026-03-20")


def appetite(path):
    body, h1, h2, small = _styles()
    s = [Paragraph("Board risk appetite statement 2026", h1),
         Paragraph("Approved by the Board on 24 February 2026. Prepared by %s, Chief Risk Officer." % P.PEOPLE["cro"], small),
         Paragraph("Credit risk", h2),
         Paragraph("The Society lends to first-time buyers, shared owners and landlords in West Yorkshire and the "
                   "Pennine towns. The Board's appetite is that accounts two or more payments in arrears stay below "
                   "1.5 per cent of accounts, that possessions stay below 0.2 per cent of accounts in any year, and that "
                   "new lending above 90 per cent loan to value stays below 20 per cent of gross lending in any year. "
                   "High loan to value lending was paused in January 2025.", body),
         Paragraph("Liquidity and capital", h2),
         Paragraph("Liquid assets of at least 14 per cent of shares and deposits; a leverage ratio of at least 5.5 per "
                   "cent; total capital of at least 16 per cent of risk-weighted assets.", body),
         Paragraph("Operational and conduct risk", h2),
         Paragraph("No tolerance for breaches of the Consumer Duty in arrears handling; complaint volumes are "
                   "reported monthly.", body)]
    _pdf(path, "Board risk appetite statement 2026", "Board risk appetite", s,
         "%s  |  Risk appetite statement 2026" % P.ORG, "2026-02-24")


def ratecard(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Product transfers"
    head = ["Product", "Loan to value up to", "Two-year fix", "Five-year fix", "Notes"]
    ws.append(head)
    rows = [["Owner-occupied", "75%", 3.80, 3.70, ""], ["Owner-occupied", "85%", 4.20, 4.10, ""],
            ["Owner-occupied", "95%", 4.60, 4.50, "above 90% current LTV"],
            ["Shared ownership", "95% of share", 4.10, 4.00, "existing borrowers only"],
            ["Buy to let", "75%", 4.50, 4.40, "interest only or capital and interest"]]
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E3E8EC")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDE", (20, 18, 13, 13, 34)):
        ws.column_dimensions[col].width = w
    for r in range(2, 7):
        for col in "CD":
            ws["%s%d" % (col, r)].number_format = '0.00"%"'
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("Terms")
    ws2.append(["Rates effective 1 October 2026 for existing borrowers switching at the end of a fixed rate. Loan to "
                "value is the current balance over the latest valuation. Rates are reviewed monthly."])
    ws2.column_dimensions["A"].width = 110
    wb.properties.creator = P.ORG
    wb.save(path)
    postprocess.scrub(path, P.ORG, "2026-09-28")
    postprocess.finish_office(path, "2026-09-28T08:12:00Z", "2026-09-28T16:47:00Z", P.ORG, person="Mortgage Pricing")


