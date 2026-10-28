"""Independent verifier: reads only target.zip (pandas, openpyxl and pypdf; none of solve.py, world.py or params.py),
recomputes every graded figure on its own code path and compares with figures.json."""
import io
import json
import os
import re
import sys
import zipfile

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from pypdf import PdfReader

HERE = os.path.dirname(os.path.abspath(__file__))
SETTINGS = {"Approved Premises": "AP", "Hostel": "HH", "Fenwold Recovery": "DS", "Pharmacy NSP": "NSP"}
CLINICS = ["HEP", "COMM", "FRC", "HHC", "SHC"]


def main(zpath):
    z = zipfile.ZipFile(zpath)
    names = z.namelist()
    fig = json.load(open(os.path.join(HERE, "figures.json")))
    n = [0]

    def check(label, cond, detail=""):
        n[0] += 1
        if not cond:
            raise AssertionError("VERIFY FAIL %s %s" % (label, detail))

    pick = lambda pre: [x for x in names if x.startswith(pre)][0]
    csv = lambda pre, **kw: pd.read_csv(io.BytesIO(z.read(pick(pre))), dtype=str, keep_default_na=False, **kw)
    text = lambda pre: re.sub(r"\s+", " ", " ".join(p.extract_text() for p in PdfReader(io.BytesIO(z.read(pick(pre)))).pages))
    r10 = lambda x: int(np.floor(x / 10 + 0.5) * 10)

    # the documents, read back through pypdf with a positive control on each
    plan, path_, intent = text("find_and_treat_van_2027_plan"), text("fenwold_hcv_pathway_2026"), text("fenwold_icb")
    for s in ("start the most people on hepatitis C treatment in 2027", "does not test again anyone it has already found",
              "within twelve weeks of the van's test", "own rates there in 2026",
              "residents of the approved premises and of the hostels on 30 September 2026",
              "each month standing for the same month of 2027", "first registrations at the pharmacies"):
        check("pdf plan: %s" % s[:40], s in plan)
    for s in ("A course of treatment ends with its last dose", "84 days (twelve weeks) or more after the last dose",
              "detected, at any level, is not a cure"):
        check("pdf pathway: %s" % s[:40], s in path_)
    check("pdf intentions: PH-12 from 1 April 2027", re.search(r"PH-12 Community pharmacy hepatitis C testing and "
                                                               r"treatment .*? 1 April 2027", intent) is not None)

    # the van's contacts: one row per contact reference, the local date from the UTC clock
    v = csv("van_tests").drop_duplicates("contact_ref")
    v["day"] = pd.to_datetime(v["recorded_at_utc"]).dt.tz_localize("UTC").dt.tz_convert("Europe/London").dt.date
    v["s"] = v["setting"].map(SETTINGS)
    check("van: every setting value mapped", v["s"].notna().all())
    reg = csv("hcv_treatment_register")
    reg["start"] = pd.to_datetime(reg["start_date"]).dt.date
    starts = reg.groupby("pseudonym")["start"].apply(list).to_dict()
    mature = pd.Timestamp("2026-07-08").date()

    def funnel(sub):
        tested = sub[sub["test_outcome"] == "Accepted"]
        found = tested[tested["rna_result"] == "Detected"]
        fm = found[found["day"] <= mature]
        st = sum(any(d <= s <= d + pd.Timedelta(days=84) for s in starts.get(p, [])) for p, d in zip(fm["client_pseudonym"], fm["day"]))
        return len(tested) / len(sub), len(found) / len(tested), st / len(fm), len(found), len(tested)
    pil = set(reg.loc[reg["prescriber_role"] == "Pharmacist independent prescriber", "prescribing_site"])
    check("pilots: three pharmacies prescribe", len(pil) == 3, pil)
    F = {s: funnel(v[v["s"] == s]) for s in ("AP", "HH", "DS", "NSP")}
    nsp = v[v["s"] == "NSP"]
    Fn, Fp = funnel(nsp[~nsp["site"].isin(pil)]), funnel(nsp[nsp["site"].isin(pil)])
    check("pilot yield", "%.3f" % Fp[1] == fig["F_pilot"] and Fp[3] == fig["pilot_found"], (Fp, fig["F_pilot"]))

    # pools
    hr, ap, ds = csv("hostel_occupancy_returns"), csv("approved_premises_returns"), csv("fenwold_recovery_new_episodes")
    ap["ym"] = pd.to_datetime(ap["Month"], format="%b-%y").dt.strftime("%Y-%m")
    in12 = lambda s: (s >= "2025-10") & (s <= "2026-09")
    hh_arr = hr.loc[in12(hr["month"]), "arrivals"].astype(int).sum()
    ap_res = ap.loc[ap["ym"] == "2026-09", "Residents at month end"].astype(int).sum()
    ap_arr = ap.loc[in12(ap["ym"]), "Arrivals"].astype(int).sum()
    ds_arr = ds.loc[in12(ds["month"]), "new_treatment_episodes"].astype(int).sum()
    nr = csv("nsp_registrations").sort_values(["registered_on", "registration_id"]).drop_duplicates("client_code")
    nr = nr[(nr["registered_on"] >= "2025-10-01") & (nr["registered_on"] <= "2026-09-30")]
    wb = load_workbook(io.BytesIO(z.read(pick("hostel_residents"))), read_only=True)
    cen = [r[0] for r in list(wb["FHT residents"].iter_rows(values_only=True))[3:] if r and r[0]]
    cen += [r[1] for r in list(wb["CSH 30.09.26"].iter_rows(values_only=True))[1:] if r and r[1]]
    hh = v[v["s"] == "HH"]
    tested = set(hh.loc[hh["test_outcome"] == "Accepted", "client_pseudonym"])
    found = set(hh.loc[hh["rna_result"] == "Detected", "client_pseudonym"])
    T = sum(p in tested for p in cen)
    Tneg = T - sum(p in found for p in cen)
    check("hostels: tested of the lists", T == fig["hh_tested"] and len(cen) == fig["census"], (T, len(cen)))
    lab = csv("virology_hcv_results")
    rv = lab[(lab["requesting_location"] == "Riverside Homeless Health") & (lab["test_code"] == "HCVAB")]
    rv = rv.drop_duplicates(["patient_pseudonym", "sample_collected", "result"]).sort_values("sample_collected")
    ev, py = 0, 0.0
    for p, g in rv.groupby("patient_pseudonym"):
        if len(g) < 2 or g["result"].iloc[0] != "Not reactive":
            continue
        d = pd.to_datetime(g["sample_collected"])
        pos = d[g["result"].values == "Reactive"]
        ev += len(pos) > 0
        py += ((pos.iloc[0] if len(pos) else d.iloc[-1]) - d.iloc[0]).days / 365.25
    check("hostels: new infections", ev == fig["inc_events"] and round(py) == fig["inc_py"], (ev, py))
    k = lambda f: f[0] * f[1]
    nr["m"] = nr["registered_on"].str[5:7].astype(int)
    q1 = nr[nr["m"] <= 3]
    found_ = {"AP": (ap_res + ap_arr) * k(F["AP"]),
              "HH": (len(cen) - T + hh_arr) * k(F["HH"]) + Tneg * F["HH"][0] * ev / py,
              "DS": ds_arr * k(F["DS"]),
              "NSP": (~q1["pharmacy_code"].isin(pil)).sum() * k(Fn) + (q1["pharmacy_code"].isin(pil).sum()
                                                                       + len(nr) - len(q1)) * k(Fp)}
    starts_ = {s: found_[s] * F[s][2] for s in found_}
    for s in found_:
        check("found %s" % s, r10(found_[s]) == fig["found10"][s] and "%.1f" % found_[s] == fig["found"][s],
              (found_[s], fig["found"][s]))
        check("started %s" % s, r10(starts_[s]) == fig["starts10"][s] and "%.1f" % starts_[s] == fig["starts"][s],
              (starts_[s], fig["starts"][s]))
    check("call", max(starts_, key=starts_.get) == fig["order"][0] == "DS")

    # the ask: last dose from the packs, a course open at the go-live followed through migrated_from
    dp = csv("daa_dispensing")
    dp["end"] = pd.to_datetime(dp["dispensed_on"]) + pd.to_timedelta(dp["days_supply"].astype(int) - 1, unit="D")
    last = dp.groupby("course_ref")["end"].max().to_dict()
    cl = csv("network_clinics")
    code = dict(zip(cl["clinic_code"], cl["clinic_code"]))
    code.update(dict(zip(cl["code_before_1_oct_2024"], cl["clinic_code"])))
    cont = reg[reg["migrated_from"] != ""].set_index("migrated_from")["course_ref"].to_dict()
    rna = lab[lab["test_code"].isin(["HCVPCR", "HCVQNT"])]
    rna = rna.assign(d=pd.to_datetime(rna["sample_collected"]))
    nd = rna[rna["result"] == "Not detected"].groupby("patient_pseudonym")["d"].apply(list).to_dict()
    mine = reg[reg["clinic_code"].isin(code) & (reg["migrated_from"] == "")].copy()
    allstarts = {}
    for p, s in zip(mine["pseudonym"], pd.to_datetime(mine["start_date"])):
        allstarts.setdefault(p, []).append(s)
    out = {}
    for _, r in mine.iterrows():
        refs = [r["course_ref"]] + ([cont[r["course_ref"]]] if r["course_ref"] in cont else [])
        ends = [last[x] for x in refs if x in last]
        if not ends:
            continue
        e = max(ends)
        s0 = pd.Timestamp(r["start_date"])
        nxt = min([s for s in allstarts[r["pseudonym"]] if s > s0], default=None)
        h = "H1" if pd.Timestamp("2024-10-01") <= e <= pd.Timestamp("2025-03-31") else (
            "H2" if pd.Timestamp("2025-04-01") <= e <= pd.Timestamp("2025-09-30") else None)
        if h and any(d >= e + pd.Timedelta(days=84) and (nxt is None or d < nxt) for d in nd.get(r["pseudonym"], [])):
            key = "%s_%s" % (code[r["clinic_code"]], h)
            out[key] = out.get(key, 0) + 1
    for c in CLINICS:
        for h in ("H1", "H2"):
            key = "%s_%s" % (c, h)
            check("cures %s" % key, out.get(key, 0) == fig["cures"][key], (out.get(key, 0), fig["cures"][key]))
    print("verify: %d assertions passed on %s" % (n[0], os.path.basename(zpath)))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(HERE), "target.zip")))
