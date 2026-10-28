"""The golden deliverables: the trustees' paper (docx, with the chart) and the funnel workbook (xlsx)."""
import io
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import params as P
import postprocess
import solve as S

DOCX = "van_2027_trustee_paper.docx"
XLSX = "van_2027_funnel.xlsx"
AUTHOR = P.DIRECTOR
SHORT = {"AP": "Approved premises", "HH": "City hostels", "DS": "Fenwold Recovery new entrants",
         "NSP": "Needle exchange new registrants"}
HALF = {1: "Oct 2024 to Mar 2025", 2: "Apr to Sep 2025"}


def chart_png(G):
    order = G["order"]
    vals = [G["starts10"][s] for s in order]
    fig, ax = plt.subplots(figsize=(6.4, 3.0), dpi=150)
    colours = ["#1F5C8B" if s == order[0] else "#A9B6C2" for s in order]
    y = range(len(order))[::-1]
    ax.barh(list(y), vals, color=colours, height=0.58)
    ax.set_yticks(list(y))
    ax.set_yticklabels([SHORT[s] for s in order], fontsize=8.5)
    for yi, v in zip(y, vals):
        ax.text(v + 0.8, yi, "%d" % v, va="center", fontsize=8.5)
    ax.text(vals[0] / 2, y[0], "Van, 2027", va="center", ha="center", color="white", fontsize=8.5, fontweight="bold")
    fig.suptitle("The needle exchange's new registrants would see the most people start treatment",
                 x=0.015, ha="left", fontsize=9.5, color="#222222")
    ax.set_xlabel("People started on hepatitis C treatment in 2027 (nearest 10)", fontsize=8.5)
    ax.set_xlim(0, max(vals) * 1.18)
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    ax.tick_params(axis="x", labelsize=8)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def _page_field(par):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    par._p.append(fld)


def docx(path, G):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10)
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.1)
    sec.top_margin = Cm(1.8)
    foot = sec.footer.paragraphs[0]
    fr = foot.add_run("Brannock Trust. Trustees' meeting, 12 November 2026. For decision. Page ")
    fr.font.size = Pt(8)
    _page_field(foot)
    o = G["order"]
    h = d.add_paragraph()
    r = h.add_run("Find-and-treat van 2027: the needle exchange's new registrants")
    r.bold = True
    r.font.size = Pt(14)
    p = d.add_paragraph("For decision. From Katie Reynolds, Director of Services, to the trustees' meeting of 12 "
                        "November 2026. This paper completes the van plan of 6 October with the figures for the "
                        "four settings.")
    p.runs[0].font.size = Pt(8.5)
    p.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    p = d.add_paragraph()
    r = p.add_run("Recommendation: the van works the needle exchange pharmacies' new registrants in 2027. It would "
                  "start about %d people on hepatitis C treatment there, %d more than at Fenwold Recovery's new "
                  "entrants, the next best at about %d." % (G["starts10"][o[0]], G["gap10"], G["starts10"][o[1]]))
    r.bold = True
    d.add_paragraph(
        "For every person the peer team found antibody positive among each setting's arrivals in the last two years, "
        "we looked at whether they still had an active infection when they were tested, from every record of their "
        "testing and treatment: the network's treatment register, Fenwold Pathology's results and Corvane Health's "
        "prison course workbook. An antibody test stays reactive after a cure, so a positive kit alone does not tell us "
        "who the van can treat. A quarter of approved premises arrivals test antibody positive, but more than half of "
        "those positives had already been cured inside prison before they arrived, and the van would start about %d "
        "people there. The needle exchange's new registrants are the people least reached by treatment so far, and "
        "most of their positives still have an active infection." % G["starts10"]["AP"])
    d.add_paragraph("What the van would find and start in 2027, by setting:")
    t = d.add_table(rows=1, cols=3)
    t.style = "Light List Accent 1"
    hdr = t.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Setting", "Active infections found", "Started on treatment"
    for s in o:
        c = t.add_row().cells
        c[0].text = SHORT[s]
        c[1].text = "%d" % G["active10"][s]
        c[2].text = "%d" % G["starts10"][s]
    for row in t.rows:
        for c in row.cells[1:]:
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    src = d.add_paragraph("People, rounded to the nearest 10; the unrounded figures are in the workbook. Sources: the "
                          "settings' arrival returns and the peer testing log, October 2024 to September 2026; Fenwold "
                          "Pathology; the network's treatment register; Corvane Health's course workbook; the 2018-19 "
                          "pilot caseload.")
    src.runs[0].font.size = Pt(8)
    src.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    d.add_picture(io.BytesIO(G["png"]), width=Cm(15.5))
    cap = d.add_paragraph("People started on treatment in 2027, by setting; the van's setting in blue.")
    cap.runs[0].font.size = Pt(8.5)
    h2 = d.add_paragraph()
    r = h2.add_run("How we sized it")
    r.bold = True
    d.add_paragraph(
        "A 2027 year is each setting's average yearly arrivals in the two years to September 2026, each person "
        "counted once in a year. The van tests %d per cent of arrivals and starts %d per cent of the people it finds "
        "with an active infection within twelve weeks, as in the 2018-19 pilot. For every arrival the peer team found "
        "antibody positive in those two years we took their status at the time of the test: cured where a course had "
        "a twelve-week test not detected and no later RNA was found, not counted where a course was still running, "
        "and otherwise the latest RNA result, including prison reception tests. Where no record says anything we used "
        "the %d per cent of first positives Fenwold Pathology finds with RNA detected."
        % (round(100 * G["test_share"]), round(100 * G["start_share"]), round(100 * G["lab_ratio"])))
    d.add_paragraph(
        "People released from prison part-way through a course are carried by the van to the end of it but, as the "
        "plan sets out, are counted neither among the infections found nor among the starts. The nurse would carry "
        "well under the 70 courses she can hold at once in any of the four settings. Clinic outcomes for the "
        "network's last full year, October 2024 to September 2025, are in the workbook.")
    d.core_properties.author = AUTHOR
    d.core_properties.title = "Find-and-treat van 2027"
    d.save(path)
    postprocess.scrub(path, P.ORG, "2026-10-21")
    postprocess.finish_office(path, "2026-10-19T09:10:00Z", "2026-10-21T15:35:00Z", P.ORG, person=AUTHOR)


def xlsx(path, G):
    wb = Workbook()
    ws = wb.active
    ws.title = "Funnel 2027"
    thin = Side(style="thin", color="B0B7BE")
    head = PatternFill("solid", fgColor="DCE6EE")
    ws.append(["Setting", "Arrivals in a year", "Tested by the van", "Antibody positive",
               "Active infection", "Started on treatment", "Active (nearest 10)", "Started (nearest 10)"])
    for s in G["order"]:
        f = G["unrounded"][s]
        ws.append([SHORT[s], round(f["arrivals"], 1), round(f["tested"], 1), round(f["positives"], 1),
                   round(f["active"], 1), round(f["starts"], 1), G["active10"][s], G["starts10"][s]])
    for row in ws.iter_rows(min_row=2, max_row=5):
        for c in row[1:6]:
            c.number_format = "#,##0.0"
        for c in row[6:8]:
            c.number_format = "0"
    ws.append([])
    ws.append(["Van tests (pilot)", G["test_share"]])
    ws.append(["Started within 12 weeks (pilot)", G["start_share"]])
    ws.append(["RNA detected, first positives (Fenwold Pathology)", G["lab_ratio"]])
    for row in ws.iter_rows(min_row=7, max_row=9):
        row[1].number_format = "0.00"
    ws2 = wb.create_sheet("Clinic outcomes")
    ws2.append(["Clinic", "Half-year (last dose)", "Confirmed cures"])
    for c in P.CLINICS:
        for h in (1, 2):
            x = G["AH"]["%s_%d" % (c, h)]
            ws2.append([P.CLINIC_LABEL[c][0].upper() + P.CLINIC_LABEL[c][1:], HALF[h], x["confirmed"]])
    ws3 = wb.create_sheet("Notes")
    notes = [
        ["Notes"],
        ["Records to 30 September 2026."],
        ["Arrivals: distinct people arriving in each year October to September, the two years to 30 September 2026 "
         "averaged (van plan, section 2)."],
        ["Tested and started: the 2018-19 pilot's shares, %d of %s arrivals tested and %d of %d people with RNA "
         "detected started within 12 weeks (van plan, section 3)." % (G["pilot_tested"], "{:,}".format(G["pilot_arrived"]),
                                                                      G["pilot_started12"], G["pilot_detected"])],
        ["Antibody positive: the peer team's reactive kits among the arrivals it tested within 60 days of arriving; "
         "invalid kits are not tests."],
        ["Active infection: each positive's status at the test, read from the network's treatment register, Corvane "
         "Health's course workbook and reception tests, and Fenwold Pathology; with no record, the share of first "
         "reactive antibody results with RNA detected."],
        ["Clinic outcomes: courses on the network's register whose last dose (the planned end, or the date a course "
         "stopped) fell in the half-year; a confirmed cure is the first RNA test 12 weeks or more after the last dose "
         "not detected (pathway, section 3), from the register or from Fenwold Pathology under every pseudonym it "
         "holds for the same patient."],
    ]
    for n in notes:
        ws3.append(n)
    ws3["A1"].font = Font(bold=True)
    ws3.column_dimensions["A"].width = 120
    for w in (ws, ws2):
        for c in w[1]:
            c.font = Font(bold=True)
            c.fill = head
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(bottom=thin)
    for col, wd in zip("ABCDEFGH", (34, 12, 12, 12, 12, 12, 12, 12)):
        ws.column_dimensions[col].width = wd
    for col, wd in zip("ABC", (30, 24, 16)):
        ws2.column_dimensions[col].width = wd
    ws.freeze_panes = "B2"
    ws2.freeze_panes = "A2"
    ws.print_title_rows = "1:1"
    ws2.print_title_rows = "1:1"
    wb.save(path)
    postprocess.scrub(path, P.ORG, "2026-10-21")
    postprocess.finish_office(path, "2026-10-14T10:02:00Z", "2026-10-21T15:12:00Z", P.ORG, person=AUTHOR)


def write(golden_dir, G):
    os.makedirs(golden_dir, exist_ok=True)
    G["png"] = chart_png(G)
    docx(os.path.join(golden_dir, DOCX), G)
    xlsx(os.path.join(golden_dir, XLSX), G)
