"""Every graded figure computed forward from the shipped files in target/ (csv and openpyxl only), with the routes the
ladder and the fork table compare. Nothing here reads world.py: what a solver can compute is what the files hold."""
import csv
import datetime as dt
import math
import os
import zoneinfo

from openpyxl import load_workbook

import params as P

D = dt.date
UK = zoneinfo.ZoneInfo("Europe/London")
MATURE = P.AS_OF - dt.timedelta(days=84)
H1, H2 = (D(2024, 10, 1), D(2025, 3, 31)), (D(2025, 4, 1), D(2025, 9, 30))
SETTING_OF = {v: k for k, v in P.TABLET_OLD.items()}


def r10(x):
    return int(math.floor(x / 10.0 + 0.5) * 10)


def edge(x, unit=10):
    return unit / 2 - abs(x - r10(x))


def rows(target, key):
    with open(os.path.join(target, P.FILES[key]), encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def iso(s):
    return D.fromisoformat(s)


# ---------------------------------------------------------------------------- the van's 2026 rates
def van(target, clock="local", since=None):
    seen, out = set(), []
    for r in rows(target, "van"):
        if r["contact_ref"] in seen:           # the tablet's double syncs
            continue
        seen.add(r["contact_ref"])
        utc = dt.datetime.fromisoformat(r["recorded_at_utc"]).replace(tzinfo=dt.timezone.utc)
        day = utc.astimezone(UK).date() if clock == "local" else utc.date()
        if since and day < since:
            continue
        out.append(dict(ref=r["contact_ref"], pid=r["client_pseudonym"], setting=SETTING_OF[r["setting"]],
                        site=r["site"], day=day, accepted=r["test_outcome"] == "Accepted",
                        found=r["rna_result"] == "Detected"))
    return out


def register(target):
    return rows(target, "register")


def pilots(reg):
    """pharmacies whose own pharmacist prescribes: the register's pharmacist-prescribed courses by site"""
    first = {}
    for r in reg:
        if r["prescriber_role"] == "Pharmacist independent prescriber":
            d = iso(r["start_date"])
            first[r["prescribing_site"]] = min(d, first.get(r["prescribing_site"], d))
    return first


def start_dates(reg):
    by = {}
    for r in reg:
        by.setdefault(r["pseudonym"], []).append(iso(r["start_date"]))
    return by


def started12(c, by):
    return any(c["day"] <= s <= c["day"] + dt.timedelta(days=84) for s in by.get(c["pid"], []))


def rates(contacts, by, pred, mature=MATURE):
    cs = [c for c in contacts if pred(c)]
    acc = [c for c in cs if c["accepted"]]
    f = [c for c in acc if c["found"]]
    fm = [c for c in f if c["day"] <= mature]
    s12 = [c for c in fm if started12(c, by)]
    return dict(offered=len(cs), tested=len(acc), found=len(f), mature=len(fm), started=len(s12),
                A=len(acc) / len(cs), F=len(f) / len(acc), ST=len(s12) / len(fm))


def km_start(contacts, by, pred):
    """Kaplan-Meier chance of a start within 84 days of the test, follow-up censored at 30 September 2026"""
    ev = []
    for c in contacts:
        if not (pred(c) and c["found"]):
            continue
        ss = [s for s in by.get(c["pid"], []) if c["day"] <= s <= c["day"] + dt.timedelta(days=84)]
        ev.append(((min(ss) - c["day"]).days, 1) if ss else (min(84, (P.AS_OF - c["day"]).days), 0))
    ev.sort()
    n, surv, i = len(ev), 1.0, 0
    while i < len(ev):
        t = ev[i][0]
        grp = [e for e in ev if e[0] == t]
        d = sum(e[1] for e in grp)
        if d:
            surv *= 1 - d / n
        n -= len(grp)
        i += len(grp)
    return 1 - surv


# ---------------------------------------------------------------------------- the 2027 pools
def census(target):
    wb = load_workbook(os.path.join(target, P.FILES["census"]), read_only=True)
    out = []
    for row in list(wb["FHT residents"].iter_rows(values_only=True))[3:]:
        if row and row[0]:
            out.append(row[0].strip())
    for row in list(wb["CSH 30.09.26"].iter_rows(values_only=True))[1:]:
        if row and row[1]:
            out.append(row[1].strip())
    return out


def in12(ym):
    return "2025-10" <= ym <= "2026-09"


def pools(target):
    hr = rows(target, "hostel_returns")
    ap = rows(target, "ap_returns")
    ds = rows(target, "ds_episodes")
    ap_month = lambda s: dt.datetime.strptime(s, "%b-%y").strftime("%Y-%m")
    regs = rows(target, "nsp_reg")
    first = {}
    for r in sorted(regs, key=lambda r: (r["registered_on"], r["registration_id"])):
        first.setdefault(r["client_code"], (iso(r["registered_on"]), r["pharmacy_code"]))
    nsp = [(d, ph) for d, ph in first.values() if D(2025, 10, 1) <= d <= P.AS_OF]
    every = [(iso(r["registered_on"]), r["pharmacy_code"]) for r in regs
             if D(2025, 10, 1) <= iso(r["registered_on"]) <= P.AS_OF]
    return dict(
        hh_arrivals=sum(int(r["arrivals"]) for r in hr if in12(r["month"])),
        ap_residents=sum(int(r["Residents at month end"]) for r in ap if ap_month(r["Month"]) == "2026-09"),
        ap_arrivals=sum(int(r["Arrivals"]) for r in ap if in12(ap_month(r["Month"]))),
        ds_arrivals=sum(int(r["new_treatment_episodes"]) for r in ds if in12(r["month"])),
        nsp=nsp, nsp_every=every)


def incidence(target):
    """new infections per person-year among residents the homeless health clinic tested and found never infected"""
    tests = {}
    for r in rows(target, "lab"):
        if r["requesting_location"] == "Riverside Homeless Health" and r["test_code"] == "HCVAB":
            tests.setdefault(r["patient_pseudonym"], set()).add((iso(r["sample_collected"]), r["result"]))
    py, ev = 0.0, 0
    for pid, ts in tests.items():
        ts = sorted(ts)
        if len(ts) < 2 or ts[0][1] != "Not reactive":
            continue
        end = next((d for d, res in ts if res == "Reactive"), None)
        ev += end is not None
        py += ((end or ts[-1][0]) - ts[0][0]).days / 365.25
    return ev / py, ev, py


# ---------------------------------------------------------------------------- the call
def solve(target, inc_scale=1.0, nsp_q1="group", transport="offer", st_method="mature", include_reissues=False,
          clock="local", since=None):
    C = van(target, clock, since)
    reg = register(target)
    by = start_dates(reg)
    pil = pilots(reg)
    PO = pools(target)
    R = {s: rates(C, by, lambda c, s=s: c["setting"] == s) for s in P.SETTINGS}
    Rn = rates(C, by, lambda c: c["setting"] == "NSP" and c["site"] not in pil)
    Rp = rates(C, by, lambda c: c["setting"] == "NSP" and c["site"] in pil)
    ST = {s: R[s]["ST"] for s in P.SETTINGS}
    if st_method == "km":
        ST = {s: km_start(C, by, lambda c, s=s: c["setting"] == s) for s in P.SETTINGS}
    cen = census(target)
    hh_t = {c["pid"] for c in C if c["setting"] == "HH" and c["accepted"]}
    hh_f = {c["pid"] for c in C if c["setting"] == "HH" and c["found"]}
    T = sum(p in hh_t for p in cen)
    Tpos = sum(p in hh_f for p in cen)
    U = len(cen) - T
    inc, ev, py = incidence(target)
    inc *= inc_scale
    nsp = PO["nsp_every"] if include_reissues else PO["nsp"]
    q1 = [(d, ph) for d, ph in nsp if d.month <= 3]
    rest = len(nsp) - len(q1)
    pq1 = sum(ph in pil for d, ph in q1)
    k = lambda x: x["A"] * x["F"]
    if nsp_q1 == "group":
        nsp_q1_found = (len(q1) - pq1) * k(Rn) + pq1 * k(Rp)
    else:
        nsp_q1_found = len(q1) * k(R["NSP"])
    post = k(Rp) if transport == "offer" else Rn["A"] * Rp["F"]
    found = {"AP": (PO["ap_residents"] + PO["ap_arrivals"]) * k(R["AP"]),
             "HH": (U + PO["hh_arrivals"]) * k(R["HH"]) + (T - Tpos) * R["HH"]["A"] * inc,
             "DS": PO["ds_arrivals"] * k(R["DS"]),
             "NSP": nsp_q1_found + rest * post}
    starts = {s: found[s] * ST[s] for s in P.SETTINGS}
    r0_found = {"AP": found["AP"], "DS": found["DS"],
                "HH": (len(cen) + PO["hh_arrivals"]) * k(R["HH"]),
                "NSP": len(nsp) * k(R["NSP"])}
    r0 = {s: r0_found[s] * ST[s] for s in P.SETTINGS}
    ladder = {"R0": dict(r0),
              "A_only": dict(r0, NSP=starts["NSP"]),
              "A_only_census_only": dict(r0, NSP=starts["NSP"], HH=len(cen) * k(R["HH"]) * ST["HH"]),
              "B_only": dict(r0, HH=starts["HH"]),
              "golden": dict(starts)}
    order = sorted(P.SETTINGS, key=lambda s: -starts[s])
    return dict(C=C, R=R, Rn=Rn, Rp=Rp, ST=ST, pilots=pil, pools=PO, census=len(cen), T=T, Tpos=Tpos, U=U,
                inc=inc, inc_events=ev, inc_py=py, nsp_regs=len(nsp), nsp_q1=len(q1), nsp_q1_pilot=pq1, nsp_rest=rest,
                found=found, starts=starts, r0=r0, r0_found=r0_found, ladder=ladder, order=order,
                found10={s: r10(found[s]) for s in P.SETTINGS}, starts10={s: r10(starts[s]) for s in P.SETTINGS})


def winner(d):
    return max(d, key=d.get)


# ---------------------------------------------------------------------------- the ask: confirmed cures by clinic
def courses(target, last_dose="dispensing", merge=True):
    """one course per treatment (a course open at the go-live is one course across both systems), its clinic, its
    last dose and its person's later course start"""
    reg = register(target)
    old2new = {c: c for c in P.CLINICS}
    old2new.update({P.CLINIC_OLD[c]: c for c in P.CLINICS})
    cont = {r["migrated_from"]: r for r in reg if r["migrated_from"]}
    packs = {}
    for r in rows(target, "dispensing"):
        packs.setdefault(r["course_ref"], []).append((iso(r["dispensed_on"]), int(r["days_supply"])))
    out = []
    for r in reg:
        if r["clinic_code"] not in old2new:
            continue
        if merge and r["migrated_from"]:
            continue
        refs = [r["course_ref"]]
        end = r["end_date"]
        if merge and r["course_ref"] in cont:
            refs.append(cont[r["course_ref"]]["course_ref"])
            end = cont[r["course_ref"]]["end_date"]
        if last_dose == "dispensing":
            pk = [p for ref in refs for p in packs.get(ref, [])]
            if not pk:
                continue
            d, n = max(pk)
            last = d + dt.timedelta(days=n - 1)
        else:
            if not end:
                continue                       # no end-of-treatment review yet
            last = iso(end)
        out.append(dict(ref=refs[0], pid=r["pseudonym"], clinic=old2new[r["clinic_code"]],
                        start=iso(r["start_date"]), last=last))
    nxt = {}
    for c in out:
        nxt.setdefault(c["pid"], []).append(c["start"])
    for c in out:
        later = [s for s in nxt[c["pid"]] if s > c["start"]]
        c["next"] = min(later) if later else None
    return out


def rna(target):
    by = {}
    for r in rows(target, "lab"):
        if r["test_code"] in ("HCVPCR", "HCVQNT"):
            by.setdefault(r["patient_pseudonym"], []).append((iso(r["sample_collected"]), r["result"]))
    return by


def cured(c, tests, cure_words=("Not detected",)):
    lo = c["last"] + dt.timedelta(days=84)
    return any(d >= lo and (c["next"] is None or d < c["next"]) and res in cure_words for d, res in tests.get(c["pid"], []))


def cures(target, last_dose="dispensing", merge=True, cure_words=("Not detected",)):
    cs = courses(target, last_dose, merge)
    t = rna(target)
    out = {}
    for clinic in P.CLINICS:
        for h, (a, b) in (("H1", H1), ("H2", H2)):
            out["%s_%s" % (clinic, h)] = sum(1 for c in cs if c["clinic"] == clinic and a <= c["last"] <= b
                                             and cured(c, t, cure_words))
    return out
