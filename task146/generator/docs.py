"""The pack's documents: the board's technical notes and its release of 8 October 2026 (PDF), the program statistics
workbook, Employer Services' notice register and the Statistics Office's concordance (xlsx), and the online job
advertisements index (JSON). The technical notes' text is kept here as data so that checks.py can quote it."""
from __future__ import annotations

import datetime as dt
import json

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import A4                                                # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import cm                                                    # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import common as K   # noqa: E402
import notices as N  # noqa: E402
import records as R  # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/"
for _n, _f in (("Carlito", "crosextra/Carlito-Regular.ttf"), ("Carlito-Bold", "crosextra/Carlito-Bold.ttf"),
               ("Carlito-Italic", "crosextra/Carlito-Italic.ttf")):
    pdfmetrics.registerFont(TTFont(_n, FONT_DIR + _f))
SANS, SANS_B, SANS_I = "Carlito", "Carlito-Bold", "Carlito-Italic"
INK = colors.HexColor("#1b2430")
GREY = colors.HexColor("#5b6573")
RULE = colors.HexColor("#9aa3ad")
PALE = colors.HexColor("#e6ecef")
BOARD = colors.HexColor("#0f4c5c")

BODY = ParagraphStyle("body", fontName=SANS, fontSize=10.0, leading=13.4, textColor=INK, spaceAfter=5)
H1 = ParagraphStyle("h1", fontName=SANS_B, fontSize=16.0, leading=20, textColor=BOARD, spaceAfter=2)
H1S = ParagraphStyle("h1s", fontName=SANS, fontSize=9.4, leading=12.2, textColor=GREY, spaceAfter=10)
H2 = ParagraphStyle("h2", fontName=SANS_B, fontSize=11.6, leading=14.5, textColor=BOARD, spaceBefore=8, spaceAfter=3)
H3 = ParagraphStyle("h3", fontName=SANS_B, fontSize=10.0, leading=13.0, textColor=INK, spaceBefore=4, spaceAfter=2)
CELL = ParagraphStyle("cell", fontName=SANS, fontSize=9.0, leading=11.0, textColor=INK)
CELLR = ParagraphStyle("cellr", parent=CELL, alignment=2)
CELLB = ParagraphStyle("cellb", fontName=SANS_B, fontSize=9.0, leading=11.0, textColor=INK)
CELLBR = ParagraphStyle("cellbr", parent=CELLB, alignment=2)
SMALL = ParagraphStyle("small", fontName=SANS, fontSize=8.4, leading=10.8, textColor=GREY, spaceBefore=3)


def _doc(path, title, author):
    return SimpleDocTemplate(path, pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=1.9 * cm,
                             bottomMargin=2.0 * cm, title=title, author=author, subject="", creator="", producer="")


def _footer(left, right):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont(SANS, 7.8)
        canvas.setFillColor(GREY)
        canvas.drawString(2.2 * cm, 1.2 * cm, left)
        canvas.drawRightString(A4[0] - 2.2 * cm, 1.2 * cm, "%s%s" % (right, doc.page))
        canvas.restoreState()
    return draw


def _table(rows, widths, right=(), bold_last=False):
    data = []
    for i, r in enumerate(rows):
        line = []
        for j, c in enumerate(r):
            hb = i == 0 or (bold_last and i == len(rows) - 1)
            st = (CELLBR if j in right else CELLB) if hb else (CELLR if j in right else CELL)
            line.append(Paragraph(str(c), st))
        data.append(line)
    t = Table(data, colWidths=widths, hAlign="LEFT")
    t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, 0), 0.6, RULE), ("LINEBELOW", (0, -1), (-1, -1), 0.6, RULE),
                           ("LINEABOVE", (0, 0), (-1, 0), 0.6, RULE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("BACKGROUND", (0, 0), (-1, 0), PALE), ("TOPPADDING", (0, 0), (-1, -1), 2.2),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2)] +
                          ([("LINEABOVE", (0, -1), (-1, -1), 0.5, RULE)] if bold_last else [])))
    return t


def n0(v):
    return "{:,}".format(int(v))


# ---------------------------------------------------------------------------------------------------- technical notes
NOTES = [
    ("1", "About these statistics", None, [
        "The Unemployment Insurance Board compiles its claims statistics from the claims system and its notice "
        "statistics from the Employer Portal. It publishes a weekly release of initial claims, monthly program "
        "statistics, open data on initial claims by industry, and the workers in layoff notices given to the board. "
        "These notes define the terms used in all four and set out how the figures are compiled. They replace the "
        "notes of March 2022."]),
    ("2", "Claims", None, []),
    ("2.1", "The claims week", "claims_week", [
        "A claims week runs from Sunday to the following Saturday. A claim takes effect on the Sunday of the week in "
        "which it is filed, and is counted in that claims week. A claims week belongs to the month in which its Sunday "
        "falls."]),
    ("2.2", "Initial claims", "initial", [
        "An initial claim is a new claim or an additional claim. A new claim is a claim to establish a benefit year. "
        "An additional claim is made by a claimant returning to an open benefit year after a spell of work. "
        "Transitional claims, made by claimants whose benefit year has ended to establish the next, are counted "
        "separately and are not initial claims. Weeks claimed by claimants already receiving benefit are not initial "
        "claims."]),
    ("2.3", "The monetary determination", "monetary", [
        "Every new claim is given a monetary determination when it is taken, unless the claimant withdraws it first. "
        "The base period is the first four of "
        "the last five completed calendar quarters before the quarter in which the claim takes effect. The claimant is "
        "monetarily eligible if earnings in the base period total at least $6,800 and were paid in at least two of "
        "its quarters. Otherwise the claimant is found monetarily ineligible and no benefit year is established. "
        "Determinations are counted in the claims week of the claim."]),
    ("2.4", "Regions", "regions", [
        "Claims are counted in the region of the claimant's home address. The eight regions have had their present "
        "boundaries since 2006. The regions file lists each region's territorial authorities and localities, with the "
        "region codes used by the Benefit Records System."]),
    ("3", "Publication", None, []),
    ("3.1", "Advance and revised counts", "advance", [
        "The weekly release, published at 08:30 on Thursday, gives the advance count of initial claims for the claims "
        "week that ended the Saturday before, nationally and by region. Each week's count is revised once, in the "
        "following week's release, to take in claims filed late for that week; the revised count is final."]),
    ("3.2", "Program statistics", "program", [
        "The program statistics for a month are published with the weekly release on the Thursday after the month's "
        "last claims week ends. They give claims, monetary determinations, first payments, weeks claimed and benefits "
        "paid, nationally and by region. A month's figures are the sum of its claims weeks. When a month is first "
        "published its last week is at the advance count, and the month is revised in the next edition. The workbook "
        "has one sheet for each fiscal year, July to June. Benefits paid are in thousands of dollars."]),
    ("3.3", "Claims systems", "systems", [
        "The UI Claims System replaced the Benefit Records System from the claims week of 27 June 2021. Weekly figures "
        "to the claims week of 20 June 2021 are kept in the Benefit Records System's export layout, and program "
        "statistics to June 2021 in its report layout, in which monetarily eligible and ineligible determinations are "
        "headed sufficient and insufficient."]),
    ("4", "Initial claims by industry", "industry", [
        "The open data give initial and transitional claims by claims week, region, industry division and claim type "
        "(N new, A additional, T transitional) from the claims week of 1 January 2023, refreshed with each weekly "
        "release. Industry is the division of the Standard Industrial Classification of the claimant's last employer, "
        "as recorded in the business register. Rows with no claims are left out. The latest week is at its advance "
        "count."]),
    ("5", "Layoff notices", "notices", [
        "An employer that plans to lay off 20 or more workers within a period of 30 days must notify the board at "
        "least 30 days before the first layoff, giving the address of each site, the number of workers to be laid off "
        "there and the date of the first layoff. Notice of a cut in the hours of 20 or more workers is given on the "
        "same form as a reduced-hours notice, which is not a layoff notice. An amended notice replaces the notice it "
        "amends. When an employer withdraws a notice, no layoffs under it are counted. Workers are counted in the "
        "region of the site at which they are to be laid off, and in the month of the notice's first layoff. Notices "
        "given before 1 July 2026 were sent to Employer Services by email and entered in its register; from 1 July "
        "2026 employers lodge them on the Employer Portal."]),
    ("6", "Enquiries", None, [
        "Nicola Pedersen, Manager, Labour Market Statistics, Unemployment Insurance Board. Telephone 04 912 6630."]),
]


def notes_text():
    return " ".join(" ".join(p) for _, _, _, p in NOTES)


def sentences(key):
    """The technical notes' sentences in the section keyed key, as they ship."""
    import re
    for _, _, k, paras in NOTES:
        if k == key:
            return [s for p in paras for s in re.split(r"(?<=[.;])\s+(?=[A-Z])", p)]
    raise KeyError(key)


def write_notes(path):
    s = [Paragraph("Unemployment Insurance Board &middot; Labour Market Statistics", H1S),
         Paragraph("Unemployment insurance claims statistics: technical notes", H1),
         Paragraph("Revised July 2026", H1S)]
    for num, head, _, paras in NOTES:
        s.append(Paragraph("%s&nbsp;&nbsp;%s" % (num, head), H2 if "." not in num else H3))
        for p in paras:
            s.append(Paragraph(p.replace("&", "&amp;"), BODY))
    doc = _doc(path, "Unemployment insurance claims statistics: technical notes", K.PEOPLE["statsmgr"])
    f = _footer("Unemployment Insurance Board  |  Claims statistics technical notes, July 2026", "page ")
    doc.build(s, onFirstPage=f, onLaterPages=f)


# ---------------------------------------------------------------------------------------------------- the release
def write_release(path, rel):
    li = K.ASOF_LAST_SUNDAY
    we = K.saturday(li)
    pe = we - dt.timedelta(days=7)
    ya = we - dt.timedelta(days=364)
    chg = rel["adv"] - rel["prev_rev"]
    word = "down" if chg < 0 else "up"
    s = [Paragraph("Unemployment Insurance Board &middot; Labour Market Statistics &middot; Weekly claims release", H1S),
         Paragraph("Initial claims: claims week ended %s" % we.strftime("%-d %B %Y"), H1),
         Paragraph("Released Thursday %s at 08:30" % K.ASOF.strftime("%-d %B %Y"), H1S),
         Paragraph("The advance count of initial claims for the claims week ended %s was %s, %s %s from the revised "
                   "count of %s for the claims week ended %s (first reported as %s). In the claims week ended %s the "
                   "board took %s initial claims." % (
                       we.strftime("%-d %B"), n0(rel["adv"]), word, n0(abs(chg)), n0(rel["prev_rev"]),
                       pe.strftime("%-d %B"), n0(rel["prev_adv"]), ya.strftime("%-d %B %Y"), n0(rel["year_ago"])), BODY),
         Spacer(1, 4)]
    head = ["Region", "Week ended %s<br/>advance" % we.strftime("%-d %b %Y"),
            "Week ended %s<br/>revised" % pe.strftime("%-d %b %Y"), "Change", "Week ended %s" % ya.strftime("%-d %b %Y")]
    rows = [head]
    for x in rel["rows"]:
        d = x["adv"] - x["prev_rev"]
        rows.append([x["region"], n0(x["adv"]), n0(x["prev_rev"]), ("+" if d > 0 else "") + n0(d) if d else "0",
                     n0(x["year_ago"])])
    rows.append(["All regions", n0(rel["adv"]), n0(rel["prev_rev"]), ("+" if chg > 0 else "") + n0(chg),
                 n0(rel["year_ago"])])
    s.append(_table(rows, [3.6 * cm, 3.3 * cm, 3.3 * cm, 2.2 * cm, 3.3 * cm], right=(1, 2, 3, 4), bold_last=True))
    s.append(Spacer(1, 6))
    s.append(Paragraph("Counts are not seasonally adjusted. The advance count is revised in the next release.", SMALL))
    s.append(Paragraph("Next release: Thursday 15 October 2026 at 08:30, claims week ended 10 October 2026.", SMALL))
    s.append(Paragraph("Definitions: Unemployment insurance claims statistics, technical notes (revised July 2026). "
                       "Enquiries: Labour Market Statistics, telephone 04 912 6630.", SMALL))
    doc = _doc(path, "Weekly claims release, claims week ended %s" % we.isoformat(), K.PEOPLE["statsmgr"])
    f = _footer("Unemployment Insurance Board  |  Weekly claims release, 8 October 2026", "page ")
    doc.build(s, onFirstPage=f, onLaterPages=f)


# ---------------------------------------------------------------------------------------------------- workbooks
LEGACY_BLOCKS = [("NEW CLAIMS", "new"), ("ADDITIONAL CLAIMS", "additional"), ("TRANSITIONAL CLAIMS", "transitional"),
                 ("MONETARY DETERMINATIONS", "determinations"), ("SUFFICIENT", "eligible"),
                 ("INSUFFICIENT", "ineligible"), ("FIRST PAYMENTS", "first_payments"),
                 ("WEEKS CLAIMED", "weeks_claimed"), ("BENEFITS PAID ($000)", "benefits")]
CURRENT_HEAD = [("New claims", "new"), ("Additional claims", "additional"), ("Transitional claims", "transitional"),
                ("Monetary determinations", "determinations"), ("Monetarily eligible", "eligible"),
                ("Monetarily ineligible", "ineligible"), ("First payments", "first_payments"),
                ("Weeks claimed", "weeks_claimed"), ("Benefits paid ($000)", "benefits")]
MON = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]


def write_program(path, T, data):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contents"
    fys = R.fiscal_years(T)
    bold = Font(bold=True)
    ws["A1"] = "Unemployment insurance program statistics"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Unemployment Insurance Board, Labour Market Statistics. Published 8 October 2026 with the September " \
               "2026 figures."
    ws["A4"] = "Sheet"
    ws["B4"] = "Months"
    ws["A4"].font = ws["B4"].font = bold
    row = 5
    for fy, months in fys.items():
        a, b = months[0], months[-1]
        ws.cell(row=row, column=1, value=fy)
        ws.cell(row=row, column=2, value="%s %d to %s %d" % (dt.date(a[0], a[1], 1).strftime("%B"), a[0],
                                                              dt.date(b[0], b[1], 1).strftime("%B"), b[0]))
        row += 1
    row += 1
    for line in ["September 2026 is provisional.",
                 "Definitions are in the technical notes (Unemployment insurance claims statistics, revised July 2026).",
                 "Enquiries: Labour Market Statistics, telephone 04 912 6630."]:
        ws.cell(row=row, column=1, value=line)
        row += 1
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 30
    head_fill = PatternFill("solid", fgColor="DCE6EA")
    order_codes = [K.LEGACY_CODE[r] for r in K.REGIONS]
    code_region = {v: k for k, v in K.LEGACY_CODE.items()}
    for fy, months in fys.items():
        ws = wb.create_sheet(fy)
        y0 = int(fy[2:6])
        allm = [(y0, m) for m in range(7, 13)] + [(y0 + 1, m) for m in range(1, 7)]
        if fy <= R.LEGACY_FY_LAST:
            ws["A1"] = "UNEMPLOYMENT INSURANCE BOARD"
            ws["A2"] = "PROGRAM STATISTICS BY REGION - FISCAL YEAR %d-%02d" % (y0, (y0 + 1) % 100)
            ws["A1"].font = ws["A2"].font = bold
            r0 = 4
            for title, key in LEGACY_BLOCKS:
                ws.cell(row=r0, column=1, value=title).font = bold
                hdr = ["REGION"] + ["%s-%02d" % (MON[m - 1], y % 100) for y, m in allm] + ["YEAR"]
                for j, h in enumerate(hdr, 1):
                    c = ws.cell(row=r0 + 1, column=j, value=h)
                    c.font = bold
                    c.alignment = Alignment(horizontal="center" if j > 1 else "left")
                for k, code in enumerate(order_codes + ["ALL"]):
                    rr = r0 + 2 + k
                    ws.cell(row=rr, column=1, value=code)
                    tot = 0
                    for j, ym in enumerate(allm, 2):
                        if ym in months:
                            v = (sum(data[(ym, x)][key] for x in K.REGIONS) if code == "ALL"
                                 else data[(ym, code_region[code])][key])
                            c = ws.cell(row=rr, column=j, value=v)
                            c.number_format = "#,##0"
                            tot += v
                    c = ws.cell(row=rr, column=14, value=tot)
                    c.number_format = "#,##0"
                    if code == "ALL":
                        for j in range(1, 15):
                            ws.cell(row=rr, column=j).font = bold
                r0 += len(order_codes) + 4
            ws.column_dimensions["A"].width = 26
            for j in range(2, 15):
                ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 9.5
        else:
            ws["A1"] = "Unemployment insurance program statistics, %d-%02d (July %d to June %d)" % (
                y0, (y0 + 1) % 100, y0, y0 + 1)
            ws["A1"].font = Font(bold=True, size=12)
            hdr = ["Month", "Region"] + [h for h, _ in CURRENT_HEAD]
            for j, h in enumerate(hdr, 1):
                c = ws.cell(row=3, column=j, value=h)
                c.font = bold
                c.fill = head_fill
                c.alignment = Alignment(horizontal="left" if j <= 2 else "right", wrap_text=True, vertical="bottom")
            rr = 4
            for ym in months:
                for reg in K.REGIONS + ["All regions"]:
                    c = ws.cell(row=rr, column=1, value=dt.date(ym[0], ym[1], 1))
                    c.number_format = "mmm yyyy"
                    c.alignment = Alignment(horizontal="left")
                    ws.cell(row=rr, column=2, value=reg)
                    for j, (_, key) in enumerate(CURRENT_HEAD, 3):
                        v = (sum(data[(ym, x)][key] for x in K.REGIONS) if reg == "All regions"
                             else data[(ym, reg)][key])
                        c = ws.cell(row=rr, column=j, value=v)
                        c.number_format = "#,##0"
                    if reg == "All regions":
                        for j in range(1, len(hdr) + 1):
                            ws.cell(row=rr, column=j).font = bold
                    rr += 1
            if fy == "FY2026-27":
                ws.cell(row=rr + 1, column=1, value="September 2026 is provisional.")
            ws.column_dimensions["A"].width = 10
            ws.column_dimensions["B"].width = 12
            for j in range(3, len(hdr) + 1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 13
            ws.row_dimensions[3].height = 30
            ws.freeze_panes = "C4"
    wb.save(path)


def write_register(path, rows):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Notices"
    ws["A1"] = "Employer Services - layoff notices received by email"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Register closed 30/06/2026. Notices from 1 July 2026 are on the Employer Portal."
    for j, h in enumerate(N.REGISTER_COLS, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFF2CC")
    for i, r in enumerate(rows, 5):
        for j, h in enumerate(N.REGISTER_COLS, 1):
            v = r[h]
            if h == "Business no.":
                v = int(v)
            c = ws.cell(row=i, column=j, value=v)
            if isinstance(v, dt.date):
                c.number_format = "dd/mm/yyyy"
    widths = [12, 11, 42, 12, 24, 15, 7, 8, 11, 13, 10, 9, 28]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "A4:%s%d" % (openpyxl.utils.get_column_letter(len(N.REGISTER_COLS)), 4 + len(rows))
    wb.save(path)


def write_concordance(path, rows):
    import openpyxl
    from openpyxl.styles import Alignment, Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Divisions"
    ws["A1"] = "Standard Industrial Classification 2023: concordance with SIC 2006 at division level"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Statistics Office, Classifications and Standards. Version 1.2, November 2025."
    hdr = ["SIC 2006 division", "SIC 2006 title", "SIC 2023 division", "SIC 2023 title", "Correspondence"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True)
    for i, r in enumerate(rows, 5):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            if j in (1, 3):
                c.number_format = "@"
    for j, w in enumerate([10, 52, 10, 48, 14], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    ws.freeze_panes = "A5"
    wb.save(path)


def write_ads(path, series):
    doc = {"dataset": "Online job advertisements index", "publisher": "Statistics Office (experimental series)",
           "frequency": "weekly, weeks ending Sunday", "base": "2019 average = 100", "unit": "index",
           "retrieved": "2026-10-07T21:14:09Z",
           "notes": "Advertisements for jobs located in each region, counted once in each week they are open, from "
                    "the main job boards. Not seasonally adjusted.",
           "observations": series}
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        json.dump(doc, fh, indent=1, ensure_ascii=False)
        fh.write("\n")
