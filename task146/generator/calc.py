"""Every figure of the main call and the asks, computed from the shipped files (or, for the analyst test, from the
clean truth through the same code): the golden and its constructions, the default path and every rival, the
precision of the golden, the leave-one-out and panel rejection margins, the twin pairs; then Ask A and Ask B from
their rows.

A view holds, for each region, the weekly initial claims by the claims week's Sunday (revised, and first reported),
the latest week at its advance count, and the monetarily ineligible determinations by claims-week month (the program
statistics), summed by quarter for the timing-blind routes and kept month by month for the golden."""
from __future__ import annotations

import csv
import datetime as dt
import math
import os
import warnings

import numpy as np
import pandas as pd
import statsmodels.api as sm

import common as K
import industry as I
import notices as N
import records as R

warnings.filterwarnings("ignore")
REG = K.REGIONS
Q3 = (2026, 3)


class View:
    pass


# ---------------------------------------------------------------------------------------------------- loading
def read_csv(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def from_pack(target):
    """Read the weekly series (both layouts), the program statistics' determinations (both layouts) and the regions
    file, as a solver would."""
    p = lambda f: os.path.join(target, f)
    code = {r["region_code"]: r["region"] for r in read_csv(p(R.F_REGIONS))}
    V = View()
    V.rev = {r: {} for r in REG}
    V.adv = {r: {} for r in REG}
    V.new = {r: {} for r in REG}
    for row in read_csv(p(R.F_LEGACY)):
        d = dt.datetime.strptime(row["WEEK_COMMENCING"], "%d/%m/%Y").date()
        r = code[row["REGION_CD"]]
        V.rev[r][d] = int(row["INITIAL_CLAIMS"])
        V.new[r][d] = int(row["NEW_CLAIMS"])
    status = {}
    for row in read_csv(p(R.F_CURRENT)):
        d = dt.date.fromisoformat(row["claims_week_start"])
        r = row["region"]
        V.rev[r][d] = int(row["initial_claims"])
        V.adv[r][d] = int(row["first_reported_initial_claims"])
        V.new[r][d] = int(row["new_claims"])
        status[d] = row["count_status"]
    V.latest = max(status)
    assert status[V.latest] == "Advance" and all(v == "Revised" for d, v in status.items() if d != V.latest)
    V.inelig_m, V.new_m = program_months(p(R.F_PROGRAM), code)
    finish(V)
    return V


def program_months(path, code):
    """Monetarily ineligible determinations and new claims by claims-week month and region, from every fiscal-year
    sheet in either layout."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    inel, new = {}, {}
    for ws in wb.worksheets:
        if not ws.title.startswith("FY"):
            continue
        rows = list(ws.iter_rows(values_only=True))
        if rows[0][0] == "UNEMPLOYMENT INSURANCE BOARD":
            block = None
            for i, r in enumerate(rows):
                if r[0] in ("INSUFFICIENT", "NEW CLAIMS"):
                    block = r[0]
                    hdr = rows[i + 1]
                    months = []
                    for h in hdr[1:13]:
                        mon, yy = h.split("-")
                        months.append((2000 + int(yy), [m.upper() for m in
                                                        ("Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split())].index(mon) + 1))
                    for rr in rows[i + 2:i + 10]:
                        reg = code[rr[0]]
                        for (ym, v) in zip(months, rr[1:13]):
                            if v is not None:
                                (inel if block == "INSUFFICIENT" else new)[(ym, reg)] = int(v)
        else:
            hdr = rows[2]
            ci, cn = hdr.index("Monetarily ineligible"), hdr.index("New claims")
            for r in rows[3:]:
                if isinstance(r[0], dt.datetime) and r[1] in REG:
                    ym = (r[0].year, r[0].month)
                    inel[(ym, r[1])] = int(r[ci])
                    new[(ym, r[1])] = int(r[cn])
    return inel, new


def from_truth(T):
    """The same view from the clean truth: every week revised but the latest, the determinations of the weeks."""
    V = View()
    V.rev = {r: {} for r in REG}
    V.adv = {r: {} for r in REG}
    V.new = {r: {} for r in REG}
    li = T.idx[K.ASOF_LAST_SUNDAY]
    for i in R.ship_index(T):
        s = T.sundays[i]
        for r in REG:
            n, a, ic = R.week_counts(T, r, i)
            V.rev[r][s] = ic
            V.new[r][s] = n
            if s >= K.CURRENT_FIRST_SUNDAY:
                V.adv[r][s] = int(T.advance[r][i])
    V.latest = K.ASOF_LAST_SUNDAY
    V.inelig_m, V.new_m = {}, {}
    for key, idx in R.months_of(T).items():
        for r in REG:
            V.inelig_m[(key, r)] = sum(int(T.adv_inelig[r][i]) if i == li else int(T.inelig[r][i]) for i in idx)
            V.new_m[(key, r)] = sum(R.week_counts(T, r, i)[0] for i in idx)
    finish(V)
    return V


def finish(V):
    V.weeks = sorted(V.rev[REG[0]])
    for r in REG:
        assert sorted(V.rev[r]) == V.weeks
    V.nat = {d: sum(V.rev[r][d] for r in REG) for d in V.weeks}
    V.pool = {}
    for (ym, r), v in V.inelig_m.items():
        q = (ym[0], (ym[1] - 1) // 3 + 1)
        V.pool[(r, q)] = V.pool.get((r, q), 0) + v
    for q in {q for _, q in V.pool}:
        V.pool[("NAT", q)] = sum(V.pool[(r, q)] for r in REG)
    V.mq3 = {r: [mpool(V, r, ym) for ym in q_months(Q3)] for r in REG + ["NAT"]}
    V.months_in_q = {}
    for (ym, r) in V.inelig_m:
        if r == REG[0]:
            q = (ym[0], (ym[1] - 1) // 3 + 1)
            V.months_in_q[q] = V.months_in_q.get(q, 0) + 1
    V.qclaims = {}
    for d in V.weeks:
        q = K.quarter_of(d)
        for r in REG:
            V.qclaims[(r, q)] = V.qclaims.get((r, q), 0) + V.rev[r][d]
            V.qclaims[("NAT", q)] = V.qclaims.get(("NAT", q), 0) + V.rev[r][d]




# ---------------------------------------------------------------------------------------------------- the turns
WCI = ("Western", "Coastal", "Island")
# the closing quarter's ineligible determinations as regressors: each claims-week month on its own (the golden), the
# quarter taken whole (timing-blind), or months grouped
SPECS = {"months": (("m1",), ("m2",), ("m3",)), "quarter": (("m1", "m2", "m3"),), "last_rest": (("m3",), ("m1", "m2")),
         "last": (("m3",),), "last_two": (("m2",), ("m3",)), "lump": (("m1",), ("m2", "m3"))}


def q_months(q):
    """The three claims-week months of a calendar quarter, in order."""
    return [(q[0], 3 * (q[1] - 1) + j) for j in (1, 2, 3)]


def mpool(V, r, ym):
    return sum(V.inelig_m[(ym, x)] for x in REG) if r == "NAT" else V.inelig_m[(ym, r)]


def terms(spec):
    return ["+".join(g) for g in SPECS[spec]]


def x26(V, r, spec):
    """The third quarter of 2026's determinations in the spec's groups of months."""
    return [sum(V.mq3[r][int(m[1]) - 1] for m in g) for g in SPECS[spec]]


def turns(V, series="rev"):
    """Every first claims week of a quarter with the closing quarter's determinations complete (three months)."""
    rows = []
    W = V.weeks
    for k in range(3, len(W)):
        d = W[k]
        q = K.quarter_of(d)
        if K.quarter_of(W[k - 1]) == q:
            continue
        pq = K.prev_quarter(q)
        if V.months_in_q.get(pq) != 3:
            continue
        row = dict(sun=d, q=q, qq=q[1], y=q[0])
        for r in REG + ["NAT"]:
            ser = (lambda dd: V.nat[dd]) if r == "NAT" else (lambda dd, r=r: V.rev[r][dd])
            row["nq_" + r] = ser(d)
            if series == "adv":
                if r == "NAT":
                    row["nqa_" + r] = sum(V.adv[x].get(d, V.rev[x][d]) for x in REG)
                else:
                    row["nqa_" + r] = V.adv[r].get(d, V.rev[r][d])
            row["p1_" + r] = ser(W[k - 1])
            row["p2_" + r] = ser(W[k - 2])
            row["p3_" + r] = ser(W[k - 3])
            row["pool_" + r] = V.pool[(r if r != "NAT" else "NAT", pq)]
            for j, ym in enumerate(q_months(pq)):
                row["m%d_%s" % (j + 1, r)] = mpool(V, r, ym)
            row["pqc_" + r] = V.qclaims[(r, pq)]
        rows.append(row)
    return pd.DataFrame(rows)


def dums(df):
    return pd.get_dummies(df["qq"].astype(str), prefix="q", dtype=float)


def design(df, r, kind="diff", base="p1", spec="months", icpt="quarter"):
    """The step of region r (or NAT) at each turn and its regressors: intercepts (one per quarter of the year, one, or
    none) and the closing quarter's determinations in the spec's groups of months (over the base week for the ratio and
    log forms)."""
    if icpt == "quarter":
        X = dums(df)
    elif icpt == "const":
        X = pd.DataFrame({"const": np.ones(len(df))}, index=df.index)
    else:
        X = pd.DataFrame(index=df.index)
    p1 = df["p1_" + r]
    for g in SPECS[spec]:
        v = sum(df["%s_%s" % (m, r)] for m in g)
        X["+".join(g)] = v / p1 if kind in ("ratio", "log") else v
    if kind == "ratio":
        y = df["nq_" + r] / p1 - 1
    elif kind == "log":
        y = np.log(df["nq_" + r] / p1)
    elif kind == "adv":
        y = df["nqa_" + r] - p1
    elif base == "mean3":
        y = df["nq_" + r] - (p1 + df["p2_" + r] + df["p3_" + r]) / 3
    else:
        y = df["nq_" + r] - p1
    return X, y


def fit_region(df, r, kind="diff", robust=False, base="p1", spec="months", icpt="quarter"):
    X, y = design(df, r, kind, base, spec, icpt)
    return sm.RLM(y, X, M=sm.robust.norms.HuberT()).fit() if robust else sm.OLS(y, X).fit()


def base_of(V, r, base="p1"):
    ser = (lambda d: V.nat[d]) if r == "NAT" else (lambda d: V.rev[r][d])
    W = V.weeks
    if base == "mean3":
        return (ser(W[-1]) + ser(W[-2]) + ser(W[-3])) / 3.0
    return float(ser(W[-1]))


def intercept(m):
    return float(m.params["q_4"]) if "q_4" in m.params else float(m.params["const"]) if "const" in m.params else 0.0


def predict(V, m, r, kind="diff", base="p1", spec="months"):
    a = intercept(m)
    lin = sum(float(m.params[t]) * x for t, x in zip(terms(spec), x26(V, r, spec)))
    b0 = base_of(V, r)
    if kind == "ratio":
        return b0 * (1 + a) + lin
    if kind == "log":
        return b0 * math.exp(a + lin / b0)
    return base_of(V, r, base) + a + lin


def regional(V, df, spec="months", icpt="quarter", **kw):
    fk = {k: v for k, v in kw.items() if k in ("kind", "robust", "base")}
    pk = {k: v for k, v in kw.items() if k in ("kind", "base")}
    return sum(predict(V, fit_region(df, r, spec=spec, icpt=icpt, **fk), r, spec=spec, **pk) for r in REG)


def national(V, df, spec="quarter", icpt="quarter", **kw):
    fk = {k: v for k, v in kw.items() if k in ("kind", "robust", "base")}
    pk = {k: v for k, v in kw.items() if k in ("kind", "base")}
    return predict(V, fit_region(df, "NAT", spec=spec, icpt=icpt, **fk), "NAT", spec=spec, **pk)


# ---------------------------------------------------------------------------------------------------- the ladder
def golden_parts(V, D, spec="months"):
    """The golden by region: the base week, the fourth-quarter intercept (the week's movement net of the refiling),
    the refiling (each month's slope times the region's determinations of that month of the third quarter), with the
    delta-method variance of each."""
    out = {}
    tn = terms(spec)
    for r in REG:
        m = fit_region(D, r, spec=spec)
        xs = np.array(x26(V, r, spec), float)
        names = ["q_4"] + tn
        Cv = m.cov_params().loc[names, names].values
        v = np.r_[1.0, xs]
        Cs = m.cov_params().loc[tn, tn].values
        sl = {t: float(m.params[t]) for t in tn}
        out[r] = dict(base=base_of(V, r), icpt=float(m.params["q_4"]), slopes=sl,
                      slope_se={t: float(m.bse[t]) for t in tn}, x=dict(zip(tn, xs)),
                      refile=float(sum(sl[t] * x for t, x in zip(tn, xs))), var_call=float(v @ Cv @ v),
                      var_refile=float(xs @ Cs @ xs), var_icpt=float(m.bse["q_4"] ** 2))
    return out


def touched_by_holiday(s):
    """A public holiday in the turn week or the week before (every January turn, and the April turns Easter reaches)."""
    return any((s + dt.timedelta(days=k)) in K.HOLIDAYS for k in range(-7, 7))


def windows(D):
    """The estimation windows every construction is tried on."""
    return {"all turns": D, "fourth-quarter turns only": D[D.qq == 4], "third and fourth-quarter turns": D[D.qq >= 3],
            "without 2020 and 2021": D[~D.y.isin([2020, 2021])], "since 2019": D[D.y >= 2019],
            "since 2017": D[D.y >= 2017], "without January turns": D[D.qq != 1],
            "the current system's turns": D[D.sun >= dt.date(2021, 7, 1)],
            "the last five years (turns from October 2021)": D[D.sun >= dt.date(2021, 10, 1)],
            "the last three years (turns from October 2023)": D[D.sun >= dt.date(2023, 10, 1)],
            "holiday-touched turns dropped": D[[not touched_by_holiday(s) for s in D.sun]]}


def ladder(V):
    D = turns(V, series="adv")
    out = {}
    cur = D[D.sun >= dt.date(2021, 7, 1)]
    q4 = D[D.qq == 4]
    Wn = windows(D)
    # the golden's family: region by region, each claims-week month of the closing quarter on its own
    G = {k: regional(V, df) for k, df in Wn.items()}
    G["Huber"] = regional(V, D, robust=True)
    G["ratio form"] = regional(V, D, kind="ratio")
    G["log-ratio form"] = regional(V, D, kind="log")
    G["mean of the last three weeks as base"] = regional(V, D, base="mean3")
    G["advance count as the target (current system)"] = regional(V, cur, kind="adv")
    G["each region with one intercept"] = regional(V, D, icpt="const")
    G["fourth-quarter turns only, one intercept"] = regional(V, q4, icpt="const")
    G["the last month and the rest of the quarter"] = regional(V, D, spec="last_rest")
    G["the last month alone"] = regional(V, D, spec="last")
    G["the last two months, the first left out"] = regional(V, D, spec="last_two")
    G["one month profile for every region, each region's own scale"] = profile(V, D)[0]
    G["the national step on the three regions' and the rest's months"] = two_groups(V, D)
    G["one intercept per quarter of the year common to every region"] = common_intercepts(V, D)
    out["golden"] = G
    gold = G["all turns"]
    out["nat_on_regional_months"] = all_months(V, D)
    out["lump"] = regional(V, D, spec="lump")
    # region-aware but timing-blind: the closing quarter's determinations taken whole
    TB = {k: regional(V, df, spec="quarter") for k, df in Wn.items()}
    TB["Huber"] = regional(V, D, spec="quarter", robust=True)
    TB["ratio form"] = regional(V, D, spec="quarter", kind="ratio")
    TB["log-ratio form"] = regional(V, D, spec="quarter", kind="log")
    TB["mean of the last three weeks as base"] = regional(V, D, spec="quarter", base="mean3")
    TB["advance count as the target (current system)"] = regional(V, cur, spec="quarter", kind="adv")
    TB["each region with one intercept"] = regional(V, D, spec="quarter", icpt="const")
    TB["each region through the origin"] = regional(V, D, spec="quarter", icpt="none")
    TB["fourth-quarter turns only, one intercept"] = regional(V, q4, spec="quarter", icpt="const")
    TB.update(region_crude(V, D))
    TB["one equation on the eight regional pools"] = all_pools(V, D)[0]
    out["timing_blind"] = TB
    # rivals
    Rv = {}
    Rv["R0 seasonal factors of the quarter's first claims week on the level of the last four weeks (default)"] = default_path(V)
    Rv["R0b ISO-week seasonal factors on the level of the last four weeks"] = iso_factor(V)
    Rv["R1 last year's week times the last eight weeks' growth on the year"] = yoy(V)
    m = sm.OLS(D["nq_NAT"] - D["p1_NAT"], dums(D)).fit()
    Rv["R2 the step from the week before, one intercept per quarter of the year"] = base_of(V, "NAT") + m.params["q_4"]
    X = dums(D)
    X["pqc"] = D["pqc_NAT"]
    m = sm.OLS(D["nq_NAT"] - D["p1_NAT"], X).fit()
    Rv["R3 the step scaled on the closing quarter's claims"] = base_of(V, "NAT") + m.params["q_4"] + \
        m.params["pqc"] * V.qclaims[("NAT", Q3)]
    tot = 0.0
    for r in REG:
        X = dums(D)
        X["pqc"] = D["pqc_" + r]
        m = sm.OLS(D["nq_" + r] - D["p1_" + r], X).fit()
        tot += base_of(V, r) + m.params["q_4"] + m.params["pqc"] * V.qclaims[(r, Q3)]
    Rv["R3r the step scaled on the closing quarter's claims, region by region"] = tot
    m = sm.OLS(D["nq_NAT"] / D["p1_NAT"] - 1, dums(D)).fit()
    Rv["R2r a quarter-start ratio"] = base_of(V, "NAT") * (1 + m.params["q_4"])
    tot = 0.0
    for r in REG:
        m = sm.OLS(D["nq_" + r] - D["p1_" + r], dums(D)).fit()
        tot += base_of(V, r) + m.params["q_4"]
    Rv["R2g quarter-start terms region by region"] = tot
    # the national determinations at the national rate (quarter taken whole), and the national months
    N4 = {k: national(V, df) for k, df in Wn.items() if k not in ("the last three years (turns from October 2023)",)}
    N4["Huber"] = national(V, D, robust=True)
    N4["ratio form"] = national(V, D, kind="ratio")
    N4["mean of the last three weeks as base"] = national(V, D, base="mean3")
    N4["one intercept"] = national(V, D, icpt="const")
    N4["through the origin"] = national(V, D, icpt="none")
    X = dums(D)
    X["pool"] = D["pool_NAT"]
    X["pqc"] = D["pqc_NAT"]
    m = sm.OLS(D["nq_NAT"] - D["p1_NAT"], X).fit()
    N4["with the closing quarter's claims"] = base_of(V, "NAT") + m.params["q_4"] + m.params["pool"] * \
        V.pool[("NAT", Q3)] + m.params["pqc"] * V.qclaims[("NAT", Q3)]
    N4["one share for every region, each region's own intercepts"] = common_share(V, D)
    out["national"] = N4
    NM = {k: national(V, df, spec="months") for k, df in Wn.items()
          if k not in ("the last three years (turns from October 2023)",)}
    NM["the last month and the rest of the quarter"] = national(V, D, spec="last_rest")
    NM["one month profile for every region, one share"] = common_share(V, D, spec="months")
    out["national_months"] = NM
    L = {}
    L4 = D.tail(4)
    L["the last four turns' step per refused claim"] = base_of(V, "NAT") + \
        float((L4["nq_NAT"] - L4["p1_NAT"]).sum() / L4["pool_NAT"].sum()) * V.pool[("NAT", Q3)]
    L1 = D.tail(1)
    L["the last turn's step per refused claim"] = base_of(V, "NAT") + \
        float((L1["nq_NAT"] - L1["p1_NAT"]).iloc[0] / L1["pool_NAT"].iloc[0]) * V.pool[("NAT", Q3)]
    ly = D[(D.y == 2025) & (D.qq == 4)]
    L["last October's step per refused claim"] = base_of(V, "NAT") + \
        float((ly["nq_NAT"] - ly["p1_NAT"]).iloc[0] / ly["pool_NAT"].iloc[0]) * V.pool[("NAT", Q3)]
    out["latest"] = L
    out["rivals"] = Rv
    out["gold"] = gold
    out["parts"] = golden_parts(V, D)
    out["parts_tb"] = golden_parts(V, D, spec="quarter")
    out["D"] = D
    return out


def profile(V, D):
    """One month profile for every region (the closing quarter's months weighted rho squared, rho and one), each region
    its own scale and quarter intercepts; rho by least squares over a grid of hundredths."""
    best = None
    for k in range(1, 151):
        rho = k / 100.0
        ssr = 0.0
        for r in REG:
            X = dums(D)
            X["w"] = rho ** 2 * D["m1_" + r] + rho * D["m2_" + r] + D["m3_" + r]
            ssr += sm.OLS(D["nq_" + r] - D["p1_" + r], X).fit().ssr
        if best is None or ssr < best[0] - 1e-9:
            best = (ssr, rho)
    rho = best[1]
    tot = 0.0
    for r in REG:
        X = dums(D)
        X["w"] = rho ** 2 * D["m1_" + r] + rho * D["m2_" + r] + D["m3_" + r]
        m = sm.OLS(D["nq_" + r] - D["p1_" + r], X).fit()
        w26 = rho ** 2 * V.mq3[r][0] + rho * V.mq3[r][1] + V.mq3[r][2]
        tot += base_of(V, r) + m.params["q_4"] + m.params["w"] * w26
    return tot, rho


def common_intercepts(V, D):
    """The regions stacked, one intercept per quarter of the year shared by every region's step, each region's own
    month slopes."""
    rows = []
    for _, t in D.iterrows():
        for r in REG:
            rows.append(dict(r=r, qq=t.qq, y=t["nq_" + r] - t["p1_" + r],
                             **{"m%d" % j: t["m%d_%s" % (j, r)] for j in (1, 2, 3)}))
    P = pd.DataFrame(rows)
    X = pd.get_dummies(P.qq.astype(str), prefix="q", dtype=float)
    for r in REG:
        for j in (1, 2, 3):
            X["m%d_%s" % (j, r)] = P["m%d" % j] * (P.r == r)
    m = sm.OLS(P.y, X).fit()
    return sum(base_of(V, r) + m.params["q_4"] + sum(m.params["m%d_%s" % (j, r)] * V.mq3[r][j - 1] for j in (1, 2, 3))
               for r in REG)


def two_groups(V, D):
    """The national step on the three regions' monthly determinations and the other five's, one intercept per quarter."""
    X = dums(D)
    x26_ = []
    for j in (1, 2, 3):
        X["hi%d" % j] = sum(D["m%d_%s" % (j, r)] for r in WCI)
        X["lo%d" % j] = D["m%d_NAT" % j] - X["hi%d" % j]
        hi = sum(V.mq3[r][j - 1] for r in WCI)
        x26_ += [("hi%d" % j, hi), ("lo%d" % j, V.mq3["NAT"][j - 1] - hi)]
    m = sm.OLS(D["nq_NAT"] - D["p1_NAT"], X).fit()
    return base_of(V, "NAT") + m.params["q_4"] + sum(m.params[k] * v for k, v in x26_)


def all_months(V, D):
    """The national step on the 24 regional monthly determinations in one equation (the regional steps discarded),
    with its standard error."""
    X = dums(D)
    cols, xv = [], []
    for r in REG:
        for j in (1, 2, 3):
            c = "m%d_%s" % (j, r)
            X[c] = D[c]
            cols.append(c)
            xv.append(V.mq3[r][j - 1])
    m = sm.OLS(D["nq_NAT"] - D["p1_NAT"], X).fit()
    v = np.r_[1.0, xv]
    names = ["q_4"] + cols
    return (base_of(V, "NAT") + m.params["q_4"] + float(m.params[cols].values @ np.array(xv)),
            math.sqrt(float(v @ m.cov_params().loc[names, names].values @ v)))


def all_pools(V, D):
    """The national step on the eight regional pools taken whole in one equation, with its standard error."""
    X = dums(D)
    for r in REG:
        X["pool_" + r] = D["pool_" + r]
    m = sm.OLS(D["nq_NAT"] - D["p1_NAT"], X).fit()
    Pv = np.array([V.pool[(r, Q3)] for r in REG])
    cols = ["pool_" + r for r in REG]
    return (base_of(V, "NAT") + m.params["q_4"] + float(m.params[cols].values @ Pv),
            math.sqrt(float(np.r_[1.0, Pv] @ m.cov_params().loc[["q_4"] + cols, ["q_4"] + cols].values @ np.r_[1.0, Pv])))


def common_share(V, D, spec="quarter"):
    """Stacked regions: region-by-quarter intercepts and one set of slopes for every region on the regions' own
    determinations (the quarter whole, or by month)."""
    rows = []
    tn = terms(spec)
    for _, t in D.iterrows():
        for r in REG:
            row = dict(r=r, qq=t.qq, y=t["nq_" + r] - t["p1_" + r])
            for g, n in zip(SPECS[spec], tn):
                row[n] = sum(t["%s_%s" % (mm, r)] for mm in g)
            rows.append(row)
    P = pd.DataFrame(rows)
    X = pd.get_dummies(P.r + "_" + P.qq.astype(str), dtype=float)
    for n in tn:
        X[n] = P[n]
    m = sm.OLS(P.y, X).fit()
    return sum(base_of(V, r) + m.params["%s_4" % r] + sum(m.params[n] * x for n, x in zip(tn, x26(V, r, spec)))
               for r in REG)


def factor_frame(V):
    w = pd.DataFrame({"sun": V.weeks, "ic": [V.nat[d] for d in V.weeks]})
    w["q"] = [K.quarter_of(s) for s in w.sun]
    kq = []
    for i in range(len(w)):
        kq.append(0 if i == 0 or w.q[i] != w.q[i - 1] else kq[-1] + 1)
    w["kq"] = kq
    w["ma"] = w.ic.rolling(52, center=True).mean()
    w["ratio"] = w.ic / w.ma
    return w


def default_path(V):
    """The seasonal route: factors of each claims week of the quarter, 2022 to 2025, on the level of the last four
    weeks, times the factor of the first claims week of the fourth quarter."""
    w = factor_frame(V)
    w2 = w[(w.sun >= dt.date(2022, 1, 2)) & (w.sun <= dt.date(2025, 12, 28))]
    fac = w2.groupby(["kq", w2.q.map(lambda x: x[1])]).ratio.mean()
    last = w.tail(4)
    lvl = np.mean([row.ic / fac[(row.kq, row.q[1])] for row in last.itertuples()])
    return float(lvl * fac[(0, 4)])


def iso_factor(V):
    w = factor_frame(V)
    w["iso"] = [(s + dt.timedelta(days=6)).isocalendar()[1] for s in w.sun]
    w2 = w[(w.sun >= dt.date(2022, 1, 2)) & (w.sun <= dt.date(2025, 12, 28))]
    fac = w2.groupby("iso").ratio.mean()
    last = w.tail(4)
    lvl = np.mean([row.ic / fac[row.iso] for row in last.itertuples()])
    return float(lvl * fac[(K.TARGET_SUNDAY + dt.timedelta(days=6)).isocalendar()[1]])


def yoy(V):
    W = V.weeks
    ly = dt.date(2025, 10, 5)
    k = W.index(ly)
    g8 = sum(V.nat[d] for d in W[-8:]) / sum(V.nat[d] for d in W[k - 8:k])
    return float(V.nat[ly] * g8)


# ---------------------------------------------------------------------------------------------------- precision
def precision(V, L):
    parts = L["parts"]
    se_call = math.sqrt(sum(p["var_call"] for p in parts.values()))
    se_refile = math.sqrt(sum(p["var_refile"] for p in parts.values()))
    se_icpt = math.sqrt(sum(p["var_icpt"] for p in parts.values()))
    se_tb = math.sqrt(sum(p["var_call"] for p in L["parts_tb"].values()))
    m = fit_region(L["D"], "NAT", spec="quarter")
    P = V.pool[("NAT", Q3)]
    names = ["q_4", "m1+m2+m3"]
    se_nat = math.sqrt(float(np.array([1.0, P]) @ m.cov_params().loc[names, names].values @ np.array([1.0, P])))
    return dict(se_call=se_call, se_refile=se_refile, se_icpt=se_icpt, se_tb=se_tb, se_nat=se_nat)


def components(V, L):
    parts = L["parts"]
    refile = sum(p["refile"] for p in parts.values())
    wci = sum(parts[r]["refile"] for r in WCI)
    without = sum(p["base"] + p["icpt"] for p in parts.values())
    pool = V.pool[("NAT", Q3)]
    pool_wci = sum(V.pool[(r, Q3)] for r in WCI)
    sep_wci = sum(V.mq3[r][2] for r in WCI)
    tb = L["parts_tb"]
    return dict(call=L["gold"], refile=refile, refile_wci_share=wci / refile, without=without, pool=pool,
                pool_wci=pool_wci, pool_wci_share=pool_wci / pool, sep_wci=sep_wci, sep_wci_share=sep_wci / pool_wci,
                slopes={r: parts[r]["slopes"] for r in REG}, base=base_of(V, "NAT"),
                tb_refile=sum(p["refile"] for p in tb.values()), tb_shares={r: tb[r]["slopes"]["m1+m2+m3"] for r in REG},
                nat_slope=float(fit_region(L["D"], "NAT", spec="quarter").params["m1+m2+m3"]))


def usual_mix(V, L):
    D = L["D"]
    hi = sum(D["pool_" + r] for r in WCI) / D["pool_NAT"]
    last = sum(D["m3_" + r] for r in WCI) / sum(D["pool_" + r] for r in WCI)
    pools = D["pool_NAT"]
    return dict(median_share=float(hi.median()), mean_share=float(hi.mean()), max_share=float(hi.max()),
                median_pool=float(pools.median()), mean_pool=float(pools.mean()), hi=hi,
                median_last=float(last.median()), max_last=float(last.max()), min_last=float(last.min()),
                q3_last=float(last[D.qq == 4].median()), last=last)


def _hat_loo(X, y):
    """Leave-one-out residuals of an OLS fit, exact, from the hat matrix."""
    X, y = np.asarray(X, float), np.asarray(y, float)
    XtXi = np.linalg.pinv(X.T @ X)
    e = y - X @ (XtXi @ (X.T @ y))
    h = np.einsum("ij,jk,ik->i", X, XtXi, X)
    return e / (1 - h)


def loo(V, D):
    """Leave one turn out, the national step's errors at each turn: the golden (each region, each month), the
    region-aware timing-blind build, months grouped, the log-ratio form, the national months, the national pool, the
    closing quarter's claims, a quarter-start term, and the 24 regional months in one equation."""
    out = {}
    yN = (D["nq_NAT"] - D["p1_NAT"]).values.astype(float)
    for name, spec in (("golden", "months"), ("timing_blind", "quarter"), ("last_rest", "last_rest"), ("lump", "lump"),
                       ("last", "last")):
        e = np.zeros(len(D))
        for r in REG:
            X, y = design(D, r, spec=spec)
            e += _hat_loo(X.values, y.values)
        out[name] = e
    e = np.zeros(len(D))
    for r in REG:
        X, y = design(D, r, kind="log", spec="months")
        p1 = D["p1_" + r].values.astype(float)
        yl = y.values.astype(float)
        e += D["nq_" + r].values - p1 * np.exp(yl - _hat_loo(X.values, yl))
    out["log_ratio"] = e
    for name, spec in (("national_months", "months"), ("pool", "quarter")):
        X, y = design(D, "NAT", spec=spec)
        out[name] = _hat_loo(X.values, y.values)
    X = dums(D)
    X["pqc"] = D["pqc_NAT"]
    out["claims"] = _hat_loo(X.values, yN)
    out["term"] = _hat_loo(dums(D).values, yN)
    X = dums(D)
    for r in REG:
        for j in (1, 2, 3):
            X["m%d_%s" % (j, r)] = D["m%d_%s" % (j, r)]
    out["all_months"] = _hat_loo(X.values, yN)
    return out


def rmse(e):
    return float(np.sqrt(np.mean(np.square(e))))


def panel(V, D):
    """Each region's step at each turn, stacked: region-by-quarter intercepts with (a) one set of month slopes for every
    region against a set per region (the regional test) and (b) one slope per region on the quarter taken whole
    against three month slopes per region (the timing test); F tests, and leave-one-turn-out errors per region and
    turn."""
    rows = []
    for k, (_, t) in enumerate(D.iterrows()):
        for r in REG:
            rows.append(dict(turn=k, r=r, qq=t.qq, y=t["nq_" + r] - t["p1_" + r], m1=t["m1_" + r], m2=t["m2_" + r],
                             m3=t["m3_" + r]))
    P = pd.DataFrame(rows)

    def X_(df, kind):
        X = pd.get_dummies(df.r + "_" + df.qq.astype(str), dtype=float)
        if kind == "common_months":
            for j in (1, 2, 3):
                X["m%d" % j] = df["m%d" % j]
        elif kind == "region_quarter":
            for r in REG:
                X["q_" + r] = (df.m1 + df.m2 + df.m3) * (df.r == r)
        else:
            for r in REG:
                for j in (1, 2, 3):
                    X["m%d_%s" % (j, r)] = df["m%d" % j] * (df.r == r)
        return X
    out = {}
    fits = {}
    for kind in ("common_months", "region_quarter", "region_months"):
        X = X_(P, kind)
        fits[kind] = sm.OLS(P.y, X).fit()
        e = []
        for tid in P.turn.unique():
            tr, te = P[P.turn != tid], P[P.turn == tid]
            Xtr = X_(tr, kind)
            Xte = X_(te, kind).reindex(columns=Xtr.columns, fill_value=0.0)
            e.extend(list(te.y - sm.OLS(tr.y, Xtr).fit().predict(Xte)))
        out["rmse_" + kind] = rmse(e)
    full = fits["region_months"]
    for kind, q in (("common_months", 21), ("region_quarter", 16)):
        out["F_" + kind] = float(((fits[kind].ssr - full.ssr) / q) / (full.ssr / full.df_resid))
        out["df_" + kind] = (q, int(full.df_resid))
    # region by region: three month slopes against one on the quarter whole, and the second month's against the last's
    out["F_region"] = {}
    out["F_lump"] = {}
    for r in REG:
        mm = fit_region(D, r)
        mq = fit_region(D, r, spec="quarter")
        ml = fit_region(D, r, spec="lump")
        out["F_region"][r] = float(((mq.ssr - mm.ssr) / 2) / (mm.ssr / mm.df_resid))
        out["F_lump"][r] = float((ml.ssr - mm.ssr) / (mm.ssr / mm.df_resid))
    out["df_region"] = (2, int(fit_region(D, REG[0]).df_resid))
    return out


def twin_timing(V, D, tol_hi=0.10, tol_nat=0.15):
    """The pair of turns of the same quarter of the year, outside 2020 and 2021, whose Western, Coastal and Island
    determinations are within tol_hi of each other and whose national determinations are within tol_nat, and whose
    three regions' last-month part differs most."""
    D = D.reset_index(drop=True).copy()
    D["step"] = D.nq_NAT - D.p1_NAT
    D["hi"] = sum(D["pool_" + r] for r in WCI)
    D["last"] = sum(D["m3_" + r] for r in WCI) / D.hi
    best = None
    for a in range(len(D)):
        for b_ in range(a + 1, len(D)):
            A, B = D.loc[a], D.loc[b_]
            if A.qq != B.qq or A.y in (2020, 2021) or B.y in (2020, 2021):
                continue
            if abs(A.hi / B.hi - 1) <= tol_hi and abs(A.pool_NAT / B.pool_NAT - 1) <= tol_nat:
                d = abs(A["last"] - B["last"])
                if best is None or d > best[0]:
                    best = (d, A, B)
    if best is None:
        return None
    _, A, B = best
    late, early = (A, B) if A["last"] > B["last"] else (B, A)
    return dict(late_turn=late.sun, early_turn=early.sun, late_hi=int(late.hi), early_hi=int(early.hi),
                late_pool=int(late.pool_NAT), early_pool=int(early.pool_NAT), late_last=float(late["last"]),
                early_last=float(early["last"]), late_step=int(late.step), early_step=int(early.step))


# ---------------------------------------------------------------------------------------------------- the asks
def ask_a_rows(target):
    import openpyxl
    p = lambda f: os.path.join(target, f)
    portal = read_csv(p(R.F_PORTAL))
    wb = openpyxl.load_workbook(p(R.F_REGISTER), data_only=True)
    ws = wb["Notices"]
    rows = list(ws.iter_rows(values_only=True))
    hi = [i for i, r in enumerate(rows) if r and r[0] == "ES ref"][0]
    hdr = rows[hi]
    register = [dict(zip(hdr, r)) for r in rows[hi + 1:] if r and r[0]]
    notes = read_csv(p(R.F_CASES))
    regions = read_csv(p(R.F_REGIONS))
    town_region = {r["locality"]: r["region"] for r in regions}
    code_region = {r["region_code"]: r["region"] for r in regions}
    return portal, register, notes, town_region, code_region


def ask_a(target, **kw):
    portal, register, notes, town_region, code_region = ask_a_rows(target)
    return N.ask_a(portal, register, notes, town_region, code_region, **kw)


def ask_b(target, **kw):
    rows = read_csv(os.path.join(target, R.F_CUBE))
    return I.ask_b(rows, **kw)


# ---------------------------------------------------------------------------------------------------- the rivals
# built on the two series the call never reads
def ads_rival(V, target):
    """The step at each turn since the job advertisements index begins, regressed on the closing quarter's change in
    the national index (the mean of the quarter's weeks over the mean of the quarter before) with one intercept per
    quarter of the year, carried to the fourth-quarter turn of 2026 on the third quarter's change."""
    import json
    with open(os.path.join(target, R.F_ADS), encoding="utf-8") as fh:
        obs = json.load(fh)["observations"]
    nat = {}
    for o in obs:
        d = dt.date.fromisoformat(o["week_ending"])
        nat.setdefault(d, []).append(o["index"])
    q_mean = {}
    for d, v in nat.items():
        q_mean.setdefault(K.quarter_of(d), []).append(float(np.mean(v)))
    q_mean = {q: float(np.mean(v)) for q, v in q_mean.items()}
    D = turns(V)
    rows = []
    for _, t in D.iterrows():
        pq = K.prev_quarter(t.q)
        ppq = K.prev_quarter(pq)
        if pq in q_mean and ppq in q_mean:
            rows.append(dict(qq=t.qq, y=t.nq_NAT - t.p1_NAT, x=q_mean[pq] / q_mean[ppq] - 1))
    P = pd.DataFrame(rows)
    X = dums(P)
    X["x"] = P.x
    m = sm.OLS(P.y, X).fit()
    x26 = q_mean[Q3] / q_mean[(2026, 2)] - 1
    return base_of(V, "NAT") + m.params["q_4"] + m.params["x"] * x26, len(P), x26




# ---------------------------------------------------------------------------------------------------- crude
# region-aware, timing-blind constructions beyond the regressions
def region_crude(V, D):
    """Crude ratios without a seasonal intercept, and the national step on two pools or on the pool and its mix."""
    b = lambda r: base_of(V, r)
    out = {}
    q4 = D[D.qq == 4]
    out["each region's mean step per refused claim at fourth-quarter turns, no intercept"] = sum(
        b(r) + float(((q4["nq_" + r] - q4["p1_" + r]) / q4["pool_" + r]).mean()) * V.pool[(r, Q3)] for r in REG)
    out["each region's summed steps over its summed pools, every turn, no intercept"] = sum(
        b(r) + float((D["nq_" + r] - D["p1_" + r]).sum() / D["pool_" + r].sum()) * V.pool[(r, Q3)] for r in REG)
    hi26 = sum(V.pool[(r, Q3)] for r in WCI)
    X = dums(D)
    X["hi"] = sum(D["pool_" + r] for r in WCI)
    X["lo"] = D["pool_NAT"] - X["hi"]
    m = sm.OLS(D["nq_NAT"] - D["p1_NAT"], X).fit()
    out["the national step on two pools (Western, Coastal and Island; the rest)"] = \
        b("NAT") + m.params["q_4"] + m.params["hi"] * hi26 + m.params["lo"] * (V.pool[("NAT", Q3)] - hi26)
    X = dums(D)
    X["pool"] = D["pool_NAT"]
    X["mix"] = sum(D["pool_" + r] for r in WCI)
    m = sm.OLS(D["nq_NAT"] - D["p1_NAT"], X).fit()
    out["the national step on the pool and its three-region part"] = \
        b("NAT") + m.params["q_4"] + m.params["pool"] * V.pool[("NAT", Q3)] + m.params["mix"] * hi26
    return out


def loo_one_share(V, D, spec="quarter"):
    """Leave one turn out: the national step from region-by-quarter intercepts and one set of slopes for every region."""
    errs = []
    tn = terms(spec)
    for k in range(len(D)):
        tr, te = D.drop(index=D.index[k]), D.iloc[[k]]
        rows = []
        for _, t in tr.iterrows():
            for r in REG:
                row = dict(r=r, qq=t.qq, y=t["nq_" + r] - t["p1_" + r])
                for g, n in zip(SPECS[spec], tn):
                    row[n] = sum(t["%s_%s" % (mm, r)] for mm in g)
                rows.append(row)
        P = pd.DataFrame(rows)
        X = pd.get_dummies(P.r + "_" + P.qq.astype(str), dtype=float)
        for n in tn:
            X[n] = P[n]
        m = sm.OLS(P.y, X).fit()
        qq = int(te.qq.iloc[0])
        p = sum(m.params["%s_%d" % (r, qq)] + sum(m.params[n] * sum(float(te["%s_%s" % (mm, r)].iloc[0]) for mm in g)
                                                 for g, n in zip(SPECS[spec], tn)) for r in REG)
        errs.append(float(te["nq_NAT"].iloc[0] - te["p1_NAT"].iloc[0] - p))
    return rmse(errs)


def twin_same_quarter(V, D, tol=0.08):
    """The pair of turns of the same quarter of the year, outside 2020 and 2021, whose national pools are within tol
    of each other and whose Western, Coastal and Island parts differ most."""
    D = D.reset_index(drop=True).copy()
    D["step"] = D.nq_NAT - D.p1_NAT
    D["hi"] = (D.pool_Western + D.pool_Coastal + D.pool_Island) / D.pool_NAT
    best = None
    for a in range(len(D)):
        for b_ in range(a + 1, len(D)):
            A, B = D.loc[a], D.loc[b_]
            if A.qq != B.qq or A.y in (2020, 2021) or B.y in (2020, 2021):
                continue
            if abs(A.pool_NAT / B.pool_NAT - 1) <= tol:
                d = abs(A.hi - B.hi)
                if best is None or d > best[0]:
                    best = (d, A, B)
    if best is None:
        return None
    _, A, B = best
    hi_, lo_ = (A, B) if A.hi > B.hi else (B, A)
    return dict(hi_turn=hi_.sun, lo_turn=lo_.sun, hi_pool=int(hi_.pool_NAT), lo_pool=int(lo_.pool_NAT),
                hi_share=float(hi_.hi), lo_share=float(lo_.hi), hi_step=int(hi_.step), lo_step=int(lo_.step))
