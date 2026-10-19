"""The pack's documents: the campaign calendar, the contact policy, the media plan, the data dictionary and the
export manifest."""
import datetime as dt
import os

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess
import records as RC

PRODUCER = P.CHAIN_DISPLAY
PDF_PRODUCER = P.CHAIN          # PDF info strings are not UTF-8


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.6, leading=13.4, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=13.5, leading=17, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=14,
                        spaceBefore=8, spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10.5, textColor=colors.HexColor("#444444"))
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=8.4, leading=10.2)
    return body, h1, h2, small, cell


def _grid(rows, widths, cell):
    t = Table([[Paragraph(str(c), cell) for c in r] for r in rows], colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA0AB")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6E9E1"))]))
    return t


def _footer(text):
    def foot(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(20 * mm, 10 * mm, text)
        canvas.drawRightString(A4[0] - 20 * mm, 10 * mm, "Page %d" % d.page)
        canvas.restoreState()
    return foot


def _doc(path, title, author, subject):
    return SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                             bottomMargin=18 * mm, title=title, author=author, creator=author, subject=subject,
                             invariant=1)


def _sections(s, sec, body, h2):
    for head, paras in sec:
        s.append(Paragraph(head, h2))
        for p in paras:
            s.append(Paragraph(p, body))


# ------------------------------------------------------------------ the contact policy (pdf)
def policy(path):
    body, h1, h2, small, cell = _styles()
    who = P.PEOPLE["dpo"]
    doc = _doc(path, "Brännvik Plus contact policy", who[0], "Marketing contact rules for Brännvik Plus members")
    s = [Paragraph("Brännvik Plus contact policy", h1),
         Paragraph("Version 2.0. Approved by the management team on 20 November 2025, applies from 1 January 2026. "
                   "Owner: %s, %s, %s." % (who[0], who[2], P.CHAIN_DISPLAY), small), Spacer(1, 6)]
    _sections(s, [
        ("1. Scope", [
            "This policy covers marketing messages to Brännvik Plus members by email, SMS, app push notification "
            "and addressed mail. Service messages (receipts, account and password messages and the welcome series on "
            "joining) fall outside the limits in section 4. The welcome series is still sent only to members who gave "
            "email consent when they joined."]),
        ("2. Legal basis and the consent record", [
            "Email, SMS and push messages are sent only with the member's consent for that channel, in line with "
            "section 19 of the Marketing Act (Marknadsföringslagen). Addressed mail is sent on legitimate interest "
            "and stops when the member objects. Consents and objections are held per member and channel in the "
            "consent platform, which went live on 1 January 2026: the state of every member at midnight was migrated "
            "from the CRM, and every change since then is logged with its time and source.",
            "A message may be sent only if the member's state for that channel is granted when it is sent."]),
        ("3. Withdrawal", [
            "Members withdraw through the unsubscribe link in every email, by replying STOPP to an SMS, in the app's "
            "settings, on My pages or through customer service. The email platform and the SMS gateway stop sending to "
            "the member at once and pass the change to the consent platform, which collects it within 15 minutes. "
            "Closing a Brännvik Plus account withdraws every channel."]),
    ], body, h2)
    s.append(Paragraph("4. Frequency limits", h2))
    s.append(_grid([["Channel", "At most", "Counted per"],
                    ["Email (marketing)", "2 messages", "member and calendar week, Monday to Sunday"],
                    ["SMS", "2 messages", "member and calendar month"],
                    ["App push", "3 messages", "member and calendar week, Monday to Sunday"],
                    ["Addressed mail", "1 piece per drop", "household"]],
                   [38 * mm, 32 * mm, 95 * mm], cell))
    _sections(s, [
        ("5. Sending hours (Swedish time)", [
            "SMS between 09:00 and 19:00 on weekdays and Saturdays, never on Sundays or public holidays. Push "
            "notifications between 08:00 and 21:00. Email at any hour; the newsletter is scheduled for 06:30."]),
        ("6. Households and addressed mail", [
            "Addressed mail goes to one member per household: the household's main card holder, meaning the member "
            "with the earliest joining date among the household's open accounts. The address file is taken on the "
            "selection date and goes to the printer the same week, so an objection logged after the selection date "
            "takes effect from the next drop."]),
        ("7. Suppression and retention", [
            "The email platform keeps every address that hard-bounced on its suppression list and does not send to it "
            "again. Send and response logs are kept for 25 months. Extracts for analysis carry member_id and "
            "household_id and never names, postal addresses, email addresses or phone numbers."]),
    ], body, h2)
    doc.build(s, onFirstPage=_footer("%s, contact policy v2.0" % P.CHAIN_DISPLAY),
              onLaterPages=_footer("%s, contact policy v2.0" % P.CHAIN_DISPLAY))
    postprocess.scrub(path, PDF_PRODUCER, "2025-11-20")


# ------------------------------------------------------------------ the media plan (pdf)
def media_plan(path):
    body, h1, h2, small, cell = _styles()
    who = P.PEOPLE["media"]
    doc = _doc(path, "Media plan 2027, Brännvik Plus", who[0], "Paid social and addressed mail, 2027")
    s = [Paragraph("Media plan 2027, Brännvik Plus", h1),
         Paragraph("Prepared by %s, %s, 10 December 2026, for %s and %s." % (
             who[0], who[2], P.PEOPLE["campaigns"][0], P.PEOPLE["crm"][0]), small), Spacer(1, 6)]
    names = {c[0]: c[1] for c in P.CAMPAIGNS}
    cpm_q = {1: 57, 2: 63, 3: 58, 4: 74}
    aud = 0.52 * 0.84 * sum(v[3] for v in P.MUNIS.values())
    rows = [["Campaign", "Flight", "Budget (SEK)", "Buying objective", "Planned reach", "Planned frequency"]]
    for camp, a, b, budget, obj in P.META:
        imp = budget / cpm_q[(a.month - 1) // 3 + 1] * 1000
        A = aud * (0.92 if camp == "BV27-14" else 1.0)
        reach = A * 0.74 * (1 - np.exp(-imp / (A * 2.1)))
        rows.append(["%s %s" % (camp, names[camp]), "%d %s to %d %s" % (a.day, a.strftime("%b"), b.day, b.strftime("%b")),
                     "{:,}".format(budget), obj, "about {:,}".format(int(round(reach, -3))),
                     "%.1f" % (imp / reach)])
    _sections(s, [("1. Paid social (Meta: Facebook and Instagram)", [
        "Every flight runs as two ad sets, one per county (Gävleborg and Västernorrland), on people aged 18 to "
        "64. The acquisition campaign Bli medlem runs all year apart from a summer pause from 3 July to 8 August and "
        "the first three days of December, and excludes existing members through a customer list refreshed monthly. "
        "Budgets are split between the counties by population. Planned reach is the estimate from Meta's planning "
        "tool at the planned CPM (Q1 57 SEK, Q2 63 SEK, Q3 58 SEK, Q4 74 SEK); it counts accounts, not people or "
        "members."])], body, h2)
    s.append(_grid(rows, [36 * mm, 30 * mm, 22 * mm, 26 * mm, 26 * mm, 22 * mm], cell))
    mrows = [["Drop", "Posting date", "Selection date", "Planned pieces", "Print and postage per piece (SEK)",
              "Planned cost (SEK)"]]
    plan_q = {"BV27-05-D1": 27500, "BV27-13-D1": 15000}
    for did, camp, posted, sel, audn, cost in P.MAIL:
        mrows.append(["%s %s" % (did, names[camp]), "%d %s" % (posted.day, posted.strftime("%B %Y")),
                      "%d %s" % (sel.day, sel.strftime("%B %Y")), "{:,}".format(plan_q[did]), "%.2f" % cost,
                      "{:,}".format(int(round(plan_q[did] * cost)))])
    _sections(s, [("2. Addressed mail", [
        "Two drops through the print house, addressed to each household's main card holder and posted as "
        "economy mail (two to three days). Vårkatalogen goes to every household with post consent; Julbrevet goes "
        "to households whose members together bought for 6,000 SEK or more from January to October 2027, with a "
        "Guld version for households with a Guld member. Undeliverable pieces come back from PostNord and are "
        "logged by the print house."])], body, h2)
    s.append(_grid(mrows, [34 * mm, 26 * mm, 26 * mm, 22 * mm, 30 * mm, 24 * mm], cell))
    _sections(s, [("3. Owned channels", [
        "SMS goes through the gateway contract at %.2f SEK per segment of up to 160 characters; a rejected message "
        "is not charged. Email and app push carry no cost per message (the platform licences sit in the CRM budget) "
        "and are planned in the campaign calendar." % P.SMS_PRICE]),
        ("4. Reporting", [
            "Meta reports reach as the number of accounts that saw an ad at least once in the reporting period of "
            "the row, so daily rows do not add up to a flight's reach. The lifetime export is taken per ad set over "
            "the whole flight. Spend is reported in SEK, excluding VAT."])], body, h2)
    doc.build(s, onFirstPage=_footer("%s, media plan 2027" % P.CHAIN_DISPLAY),
              onLaterPages=_footer("%s, media plan 2027" % P.CHAIN_DISPLAY))
    postprocess.scrub(path, PDF_PRODUCER, "2026-12-10")


# ------------------------------------------------------------------ the campaign calendar (xlsx)
def _head(ws, row, head, widths):
    for j, h in enumerate(head, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E2E8D8")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for k, w_ in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(k)].width = w_


def calendar(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Kampanjer 2027"
    ws.cell(row=1, column=1, value="Brännvik Plus campaign calendar 2027").font = Font(bold=True, size=12)
    ws.cell(row=2, column=1, value="Kept by %s, %s. Budgets are planned media and production in thousand SEK, "
                                   "excluding VAT." % (P.PEOPLE["campaigns"][0], P.PEOPLE["campaigns"][2]))
    _head(ws, 4, ["Code", "Campaign", "Purpose", "Channels", "Start", "End", "Audience", "Budget (kSEK)", "Owner"],
          [10, 26, 12, 30, 11, 11, 62, 13, 18])
    r = 5
    for code, name, purpose, chans, a, b, aud, budget, owner in P.CAMPAIGNS:
        vals = [code, name, purpose, chans, dt.date.fromisoformat(a), dt.date.fromisoformat(b), aud,
                budget if budget else None, P.PEOPLE[owner][0]]
        for j, x in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=x)
            if isinstance(x, dt.date):
                c.number_format = "yyyy-mm-dd"
            if j == 7:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    ws.freeze_panes = "A5"

    ws2 = wb.create_sheet("Utskick")
    _head(ws2, 1, ["Drop", "Campaign", "Channel", "Planned (Swedish time)", "Audience", "Subject, title or text"],
          [14, 10, 13, 20, 40, 52])
    drops = []
    for k, x in enumerate(P.NEWSLETTER):
        drops.append(("BV27-01-N%02d" % (k + 1), "BV27-01", "email", dt.datetime(x.year, x.month, x.day, 6, 30),
                      "Members with email consent", "Plusbrevet: " + P.THEMES[k]))
    for mid, camp, tt, aud, subj, _ in P.EMAILS:
        drops.append((mid, camp, "email", dt.datetime(*tt), _aud(aud), subj))
    for sid, camp, tt, aud, chars in P.SMS:
        drops.append((sid, camp, "sms", dt.datetime(*tt), _aud(aud), "SMS text, %d characters" % chars))
    for pid, camp, tt, aud, title in P.PUSH:
        drops.append((pid, camp, "push", dt.datetime(*tt), _aud(aud), title))
    for did, camp, posted, sel, aud, cost in P.MAIL:
        drops.append((did, camp, "direct mail", dt.datetime(posted.year, posted.month, posted.day, 0, 0),
                      "Households with post consent" if aud == "households_post" else
                      "Households buying for 6,000 SEK or more, January to October 2027",
                      "Posted %s, selection %s" % (posted.isoformat(), sel.isoformat())))
    drops.sort(key=lambda z: (z[3], z[0]))
    for s_, (wid, off, subj) in enumerate(P.WELCOME):
        drops.append((wid, "BV27-02", "email", None, "Automation: new members with email consent, %s" % (
            "on joining" if off == 0 else "%d days after joining at 07:15" % off), subj))
    for k, row in enumerate(drops, 2):
        for j, x in enumerate(row, 1):
            c = ws2.cell(row=k, column=j, value=x)
            if isinstance(x, dt.datetime):
                c.number_format = "yyyy-mm-dd" if row[2] == "direct mail" else "yyyy-mm-dd hh:mm"
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("Erbjudanden")
    _head(ws3, 1, ["Offer code", "Campaign", "Offer", "Valid from", "Valid to", "Loaded to"], [13, 12, 44, 11, 11, 58])
    rows = []
    for code, camp, text, kind, val, base in P.OFFERS:
        c = [z for z in P.CAMPAIGNS if z[0] == camp][0]
        if camp == "BV27-02":
            rows.append((code, camp, text, None, None, "Every member joining in 2027, valid 30 days from joining"))
            continue
        loaded = {"BV27-10": "Members with purchases January to May 2027 and none June to August 2027",
                  "BV27-12": "Guld members (2027 tier)"}.get(camp, "Every open account at the start of the campaign")
        rows.append((code, camp, text, dt.date.fromisoformat(c[4]), dt.date.fromisoformat(c[5]), loaded))
    for k, x in enumerate(P.NEWSLETTER):
        kind, val = P.NEWS_OFFER[k]
        text = ("%d kr rabatt på %s" % (val, P.NEWS_ITEMS[k])) if kind == "fixed" else \
            ("%d %% på %s" % (val, P.NEWS_ITEMS[k]))
        rows.append(("BV27-01-O%02d" % (k + 1), "BV27-01", text, x, x + dt.timedelta(days=13),
                     "Every open account on the issue date"))
    rows.sort(key=lambda z: z[0])
    for k, row in enumerate(rows, 2):
        for j, x in enumerate(row, 1):
            c = ws3.cell(row=k, column=j, value=x)
            if isinstance(x, dt.date):
                c.number_format = "yyyy-mm-dd"
    ws3.freeze_panes = "A2"
    wb.save(path)
    postprocess.scrub(path, PRODUCER, "2027-12-28")
    postprocess.finish_office(path, "2026-11-30T09:12:00Z", "2027-12-28T14:41:00Z", PRODUCER)


def _aud(a):
    return {"all": "All members with consent for the channel", "county:X": "Members in Gävleborg with SMS consent",
            "lapsed": "Members with purchases January to May 2027 and none June to August 2027",
            "guld": "Guld members (2027 tier)"}[a]


# ------------------------------------------------------------------ the data dictionary (md) and the manifest (csv)
def dictionary(path):
    F = RC.FILES
    crm = P.PEOPLE["crm"]
    text = f"""# Br\u00e4nnvik Plus 2027 reach extracts: data dictionary

Prepared by the CRM team ({crm[0]}, {crm[2]}), 12 January 2028. One section per file in this folder; row counts and
extract times are in `{F["manifest"]}`. Member-level files carry `member_id` and `household_id` only. Names, postal
addresses, email addresses and phone numbers stay in the source systems (contact policy, section 7).

## Conventions

- `member_id` is `BP` followed by eight digits, the last of them a check digit (Luhn). It is the member key in every
  member-level file; the SMS gateway calls it `customer_ref` and the direct mail file `addressee_member_id`.
- `household_id` groups the accounts registered at the same address.
- Each system writes time its own way, stated per file: the consent platform and the email platform in UTC, the
  SMS gateway and the POS in Swedish local time without an offset, the app messaging platform in ISO 8601 with the
  Swedish offset, Meta as dates in the ad account's time zone (Europe/Stockholm). Sweden moved to summer time on
  28 March 2027 and back on 31 October 2027.
- Amounts are SEK. Till amounts include VAT; Meta spend excludes VAT.
- Campaign codes (`BV27-01` to `BV27-14`) and drop and offer codes are those of `{F["calendar"]}`.

## {F["members"]}

CRM extract taken on 8 January 2028. One row per Br\u00e4nnvik Plus account that was open at any time on or after
1 January 2026; accounts closed before that date are not in the extract.

| Column | Meaning |
|---|---|
| member_id | the account |
| card_number | the member's current card, 16 digits |
| household_id | the household the account is registered to |
| joined_date | the date the account was opened |
| signup_channel | where the account was opened: store_kiosk, web_signup or app_signup |
| home_store | the member's home store (`store_code`), chosen at signup and editable on My pages |
| postcode_area | the first three digits of the member's postcode |
| municipality | the municipality of the member's address |
| birth_year | year of birth (membership is from age 16) |
| email_on_file, mobile_on_file | Y when the CRM holds an email address or a mobile number |
| app_registered_date | the date the member first signed in to the app; blank if never |
| tier_2027 | the tier for 2027, set on 1 January 2027 from purchases paid in 2026; blank for accounts closed in 2026 |
| tier_2028 | the tier for 2028, set on 1 January 2028 from purchases paid in 2027; blank for closed accounts |
| account_status | open or closed on the extract date |
| closed_date | the date the account was closed |

Tiers: Silver from {P.SILVER:,} SEK and Guld from {P.GULD:,} SEK of purchases paid with the member's card in the
calendar year, Bas below that.

## {F["snapshot"]}

The consent state migrated into the consent platform when it went live on 1 January 2026 at 00:00 Swedish time, one
row per member who had joined by then. `email`, `sms`, `push` and `post` are Y (granted) or N. `post` is addressed
mail, where N means the member has objected.

## {F["consent"]}

The consent platform's log: one row per change in a member's state on one channel, from 1 January 2026 to
31 December 2027.

| Column | Meaning |
|---|---|
| event_id | the log's running number |
| member_id | the account |
| channel | email, sms, push or post |
| action | granted or withdrawn |
| event_time_utc | when the platform recorded the change (UTC) |
| source | store_kiosk, web_signup or app_signup (given on joining); my_pages; customer_service; app_settings; esp_unsubscribe and esp_complaint (from the email platform); sms_stop (STOPP replies, from the SMS gateway); account_closed |

A member's state at a moment is the snapshot state (N for members who joined later) updated by the log's rows up to
that moment. Changes made at the email platform or the SMS gateway reach the log with the next collection, up to 15
minutes after the member acted; both systems stop sending to the member at once.

## {F["calendar"]}

The campaign manager's calendar ({P.PEOPLE["campaigns"][0]}). `Kampanjer 2027` has one row per campaign (code, name,
purpose, channels, dates, audience, planned budget in thousand SEK, owner). `Utskick` has one row per planned drop:
email mailings, SMS and push sends with their planned Swedish time, and addressed mail drops with their posting and
selection dates; the welcome automation is listed at the end. `Erbjudanden` has one row per personal offer loaded
to members' cards, with its validity and the members it was loaded to.

## {F["mailings"]}

Email platform: one row per mailing. `mailing_id` is BV27-01-N01 to N26 for the newsletter issues, BV27-xx-E1 and
E2 for campaign mailings and BV27-02-W1 to W3 for the welcome automation. `type` is newsletter, campaign or
automation; `scheduled_utc` is blank for the automation; `sends` is the mailing's rows in `{F["sends"]}`.

## {F["sends"]}

Email platform: one row per message sent to a member in 2027.

| Column | Meaning |
|---|---|
| send_id | the message |
| mailing_id | the mailing |
| member_id | the subscriber key |
| sent_at | when the message left the platform (UTC) |
| status | delivered; soft_bounce (temporary failure, not retried); hard_bounce (the address does not exist; the platform suppresses it from then on) |

## {F["events"]}

Email platform: one row per tracked response to a message. A message can have several opens.

| Column | Meaning |
|---|---|
| event_id, send_id, mailing_id, member_id | the response and the message it belongs to |
| event_type | open, click, unsubscribe or spam_complaint |
| event_at | when the platform recorded it (UTC) |
| client | on opens, the mail client that loaded the tracking image. Apple Mail Privacy Protection marks images loaded through Apple's proxy, which Apple Mail does for most delivered messages whether or not they are read |
| link | on clicks: offer, recipes, store_finder, app_download or other |

## {F["sms"]}

SMS gateway: one row per message submitted.

| Column | Meaning |
|---|---|
| message_id | the gateway's id |
| sender | the alphanumeric sender shown on the phone |
| campaign_ref | the drop (`Utskick`) |
| customer_ref | the member_id |
| operator | the recipient's network |
| submitted, status_time | Swedish local time, no offset |
| status | the operator's delivery receipt: DELIVRD delivered to the handset; UNDELIV undeliverable; EXPIRED not delivered within the 48-hour validity period; REJECTD refused by the network and not charged |
| segments | message parts of up to 160 characters |
| price_sek | the gateway's charge |
| reply_text, reply_time | STOPP replies, which the gateway passes on as withdrawals |

## {F["push"]}

App messaging platform: one JSON object per line, one per notification sent to a member's device. Keys:
`notification_id` (the drop), `member_id`, `platform` (ios or android), `sent_at` (ISO 8601 with the Swedish
offset), `status` (delivered or failed), `failure_reason` (on failed rows: unregistered when the app is no longer
installed, expired_token when the device token had lapsed) and `opened_at` (only when the member opened the app
from the notification).

## {F["mail"]}

The print house's address file with PostNord's returns: one row per piece. `addressee_member_id` is the household's
main card holder on the selection date; `version` is Standard or Guld; `returned_date` is the date PostNord
returned the piece as undeliverable.

## {F["meta_daily"]} and {F["meta_life"]}

Meta Ads Manager exports, downloaded by {P.PEOPLE["media"][0]}. The daily file has one row per ad set and day with
delivery; the lifetime file one row per ad set over its flight. Reach counts accounts that saw an ad at least once in
the row's reporting period, so daily reach does not add up to lifetime reach. Frequency is Impressions over Reach.

## {F["spend"]}

BI warehouse: one row per member, store and month with at least one purchase with the member's card in 2027.

| Column | Meaning |
|---|---|
| member_id, store_code | the member and the store |
| month | YYYY-MM |
| visits | receipts with the member's card |
| spend_sek | amount paid, after offer discounts |
| offer_discount_sek | discounts from personal offers on those receipts |
| offers_redeemed | personal offers redeemed on those receipts |

## {F["redemptions"]}

POS loyalty engine: one row per personal offer redeemed at a till, semicolon-separated with decimal commas as the
POS exports it. `receipt_id` is the store and the receipt number, `redeemed_at` is Swedish local time,
`receipt_total_sek` is the amount paid after the discount. A receipt redeems one offer. These receipts are part of
`{F["spend"]}`.

## {F["stores"]} and {F["postcodes"]}

`stores.csv`: one row per store (code, name, municipality, county, postcode area, format Stor, Mellan or N\u00e4ra,
sales area in square metres, year opened). `postcode_areas.csv`: one row per three-digit postcode area in the two
counties, with its municipality and municipality code, county, population on 31 December 2026 and households.
"""
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(text)


MANIFEST = [
    ("members", "CRM", "2028-01-08 06:12", "dates"),
    ("snapshot", "Consent platform", "2028-01-08 06:20", "state at 2026-01-01 00:00 Swedish time"),
    ("consent", "Consent platform", "2028-01-08 06:21", "UTC"),
    ("calendar", "Campaign team workbook", "2028-01-10 09:05", "Swedish time"),
    ("mailings", "Email platform", "2028-01-09 10:41", "UTC"),
    ("sends", "Email platform", "2028-01-09 10:44", "UTC"),
    ("events", "Email platform", "2028-01-09 10:52", "UTC"),
    ("sms", "SMS gateway", "2028-01-09 11:05", "Swedish local time, no offset"),
    ("push", "App messaging platform", "2028-01-09 11:20", "ISO 8601 with offset"),
    ("mail", "Print house", "2028-01-10 13:37", "dates"),
    ("meta_daily", "Meta Ads Manager", "2028-01-10 15:02", "dates, Europe/Stockholm"),
    ("meta_life", "Meta Ads Manager", "2028-01-10 15:04", "dates, Europe/Stockholm"),
    ("spend", "BI warehouse", "2028-01-08 07:30", "months"),
    ("redemptions", "POS loyalty engine", "2028-01-08 07:48", "Swedish local time, no offset"),
    ("stores", "CRM", "2028-01-08 06:12", ""),
    ("postcodes", "BI warehouse", "2028-01-08 07:31", ""),
    ("policy", "Data protection officer", "2025-11-20", ""),
    ("mediaplan", "Media buyer", "2026-12-10", ""),
    ("dictionary", "CRM team", "2028-01-12", ""),
]


def manifest(path, rows):
    import csv
    with open(path, "w", encoding="utf-8", newline="") as fh:
        wr = csv.writer(fh, lineterminator="\n")
        wr.writerow(["file", "source_system", "extracted", "rows", "time_basis"])
        for k, system, when, basis in MANIFEST:
            wr.writerow([RC.FILES[k], system, when, rows.get(k, ""), basis])
