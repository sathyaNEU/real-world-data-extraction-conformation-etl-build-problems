"""task157 (architecture 2) shipped documents: the office's filing and sample record for serial I-18-2028,
the verification instruction the county recorders work to, Navajo County's own 2019 disposition code list,
and the 2024 full count's close-out workbook.

Their text is kept here as data so that `checks.py` can quote and execute it, and every absence assertion in
`checks.py` runs with a positive control through the same extractor."""
from __future__ import annotations

import datetime as dt
import os
import sys

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                        # noqa: E402
from reportlab.lib.pagesizes import letter                                              # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                         # noqa: E402
from reportlab.lib.units import inch                                                    # noqa: E402
from reportlab.pdfbase import pdfmetrics                                                # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                            # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle   # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import common as K          # noqa: E402
import postprocess as PP    # noqa: E402
import records as R         # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/"
for _n, _f in (("Carlito", "crosextra/Carlito-Regular.ttf"), ("Carlito-Bold", "crosextra/Carlito-Bold.ttf"),
               ("LibSans", "liberation/LiberationSans-Regular.ttf"),
               ("LibSans-Bold", "liberation/LiberationSans-Bold.ttf")):
    pdfmetrics.registerFont(TTFont(_n, FONT_DIR + _f))

INK = colors.HexColor("#1b1f26")
GREY = colors.HexColor("#5a6270")
SEAL = colors.HexColor("#1f3b63")
RULE = colors.HexColor("#aab2bd")
HEAD = colors.HexColor("#e6eaf0")

F_RECORD = "sos_filing_and_selection_record_i18-2028.pdf"
F_INSTRUCTION = "county_recorder_verification_instruction_2027.pdf"
F_OLDCODES = "navajo_recorder_disposition_codes_2019.txt"


def styles(font="Carlito", bold="Carlito-Bold", accent=SEAL, size=10.1):
    return dict(
        body=ParagraphStyle("b", fontName=font, fontSize=size, leading=size * 1.38, textColor=INK,
                            spaceAfter=5.0),
        h1=ParagraphStyle("h1", fontName=bold, fontSize=14.4, leading=18.0, textColor=accent, spaceAfter=3.0),
        sub=ParagraphStyle("s", fontName=font, fontSize=8.8, leading=11.6, textColor=GREY, spaceAfter=9.0),
        h2=ParagraphStyle("h2", fontName=bold, fontSize=10.6, leading=13.8, textColor=accent, spaceBefore=7.0,
                          spaceAfter=2.5),
        cell=ParagraphStyle("c", fontName=font, fontSize=8.5, leading=10.5, textColor=INK),
        cellr=ParagraphStyle("cr", fontName=font, fontSize=8.5, leading=10.5, textColor=INK,
                             alignment=2),
        cellb=ParagraphStyle("cb", fontName=bold, fontSize=8.5, leading=10.5, textColor=INK),
        small=ParagraphStyle("sm", fontName=font, fontSize=8.1, leading=10.4, textColor=GREY, spaceBefore=4.0),
    )


class _Doc(SimpleDocTemplate):
    """The office's own documents carry the office as producer and creator: a state record whose producer
    string names a writing library is a defect (reduce-house-fixes H1)."""

    def __init__(self, path, title, author, **kw):
        self._meta = (title, author)
        SimpleDocTemplate.__init__(self, path, pagesize=letter, title=title, author=author, subject="",
                                   creator=author, leftMargin=1.0 * inch, rightMargin=1.0 * inch,
                                   topMargin=0.9 * inch, bottomMargin=0.85 * inch, **kw)

    def handle_documentBegin(self):
        SimpleDocTemplate.handle_documentBegin(self)
        title, author = self._meta
        c = self.canv
        c.setTitle(title)
        c.setAuthor(author)
        c.setSubject("")
        c.setCreator(author)
        c.setProducer("%s document services" % K.OFFICE)


def _doc(path, title, author):
    return _Doc(path, title, author)


def _footer(left, font="Carlito"):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 7.5)
        canvas.setFillColor(GREY)
        canvas.drawString(1.0 * inch, 0.5 * inch, left)
        canvas.drawRightString(letter[0] - 1.0 * inch, 0.5 * inch, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _table(st, head, rows, widths, align_right=()):
    data = [[Paragraph(esc(h), st["cellb"]) for h in head]]
    for r in rows:
        data.append([Paragraph(esc(str(v)), st["cellr"] if i in align_right else st["cell"])
                     for i, v in enumerate(r)])
    t = Table(data, colWidths=widths, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), HEAD),
        ("LINEBELOW", (0, 0), (-1, 0), 0.6, RULE),
        ("LINEBELOW", (0, 1), (-1, -2), 0.25, colors.HexColor("#dfe3ea")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3.0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3.0),
        ("LEFTPADDING", (0, 0), (-1, -1), 4.0),
    ]))
    return t


# ---------------------------------------------------------------------------------- the filing record
def record_text():
    """The office's record of the filing and of the sample it drew. Every figure in it is the office's own
    count or the statute's own number, and it is the highest authority in the pack."""
    lines = "{:,}".format(K.LINES_FILED)
    return [
        ("Measure", [
            "Serial number %s, a measure proposing an amendment to the Constitution of Arizona, filed by "
            "the Arizona Water Security Committee, a political committee registered with this office."
            % K.SERIAL,
            "Petition sheets for serial %s were issued to the committee on %s. The committee filed the "
            "petition at this office on %s." % (K.SERIAL, K.long_date(K.SERIAL_ISSUED), K.long_date(K.FILED)),
        ]),
        ("Receipt", [
            "The committee presented %s petition sheets. %d of those sheets were sheets already filed "
            "with this office under the same serial number and were struck at receipt; they are not part of "
            "the filing. The office receipted %s sheets, bearing %s signature lines, and the receipt was "
            "closed on %s." % ("{:,}".format(K.SHEETS_PRESENTED), K.SHEETS_STRUCK,
                               "{:,}".format(K.SHEETS), lines, K.long_date(K.RECEIPT_CLOSED)),
            "Each receipted sheet carries the registration number of the circulator who gathered it and the "
            "number printed on the sheet. Sheets are issued to a circulator in a numbered pad, so that "
            "number is the circulator's own series and runs from one.",
        ]),
        ("Circulator registrations", [
            "A paid circulator is registered with this office by the firm that engages it, under that firm's "
            "own registration number. A circulator who leaves one firm for another registers again under the "
            "firm it joined: the earlier registration is withdrawn as of the day the new registration takes "
            "effect, and sheets gathered on that day are gathered under the new registration.",
            "A firm that changes its trading name keeps its registration number. The name shown on a "
            "registration is the name the firm was trading under on the day the registration was taken.",
        ]),
        ("Signatures required", [
            "A measure proposing a constitutional amendment requires signatures equal to fifteen per cent of "
            "the total votes cast for Governor at the last general election at which a Governor was elected. "
            "%s votes were cast for Governor at the general election of %s, and the number required for "
            "serial %s is therefore %s." % ("{:,}".format(K.GOV_VOTES_2026), K.long_date(K.ELECTION_2026),
                                            K.SERIAL, "{:,}".format(K.REQUIRED)),
        ]),
        ("The random sample", [
            "The office drew a random sample of five per cent of the signature lines filed, %s of the %s "
            "lines receipted, on %s. The selection was made by the office's random selection program over "
            "the receipted sheets and lines." % ("{:,}".format(K.SAMPLE_LINES), lines,
                                                 K.long_date(K.SAMPLE_DRAWN)),
            "The office transmitted to each county recorder a copy of every sampled signature line assigned "
            "to that recorder, with the sheet and the line it was taken from. A line is assigned to the "
            "recorder of the county shown by the residence address written on it.",
            "Each recorder shall complete the verification and return its determinations within fifteen "
            "business days of transmittal, which for this filing closes on %s."
            % K.long_date(K.RETURNS_CLOSE),
        ]),
        ("Determination", [
            "If the number of valid signatures established for the petition is less than ninety-five per "
            "cent of the number required, %s signatures, the petition is insufficient. If it is more than "
            "one hundred and five per cent of the number required, %s signatures, the petition is "
            "sufficient. If it is neither, the Secretary shall order the examination of all signature lines "
            "filed." % ("{:,}".format(K.BAND_LO), "{:,}".format(K.BAND_HI)),
            "The Secretary certifies the measures for the general election of %s no later than %s."
            % (K.long_date(K.ELECTION_2028), K.long_date(K.CERTIFY_BY)),
        ]),
    ]


def write_record(path):
    st = styles()
    s = [Paragraph("Filing and sample selection record: initiative serial %s" % K.SERIAL, st["h1"]),
         Paragraph("%s, %s. Record opened %s. Prepared by %s, State Elections Director."
                   % (K.OFFICE, K.DIVISION, K.long_date(K.SAMPLE_DRAWN), K.PEOPLE["director"]), st["sub"])]
    for head, paras in record_text():
        s.append(Paragraph(esc(head), st["h2"]))
        for pr in paras:
            s.append(Paragraph(esc(pr), st["body"]))
    s.append(Spacer(1, 6))
    s.append(_table(st, ["Item", "Figure"], [
        ["Petition sheets presented", "{:,}".format(K.SHEETS_PRESENTED)],
        ["Petition sheets struck at receipt", "{:,}".format(K.SHEETS_STRUCK)],
        ["Petition sheets receipted", "{:,}".format(K.SHEETS)],
        ["Signature lines filed", "{:,}".format(K.LINES_FILED)],
        ["Signature lines drawn into the sample", "{:,}".format(K.SAMPLE_LINES)],
        ["Signatures required", "{:,}".format(K.REQUIRED)],
        ["Ninety-five per cent of the number required", "{:,}".format(K.BAND_LO)],
        ["One hundred and five per cent of the number required", "{:,}".format(K.BAND_HI)],
    ], [3.5 * inch, 1.6 * inch], align_right=(1,)))
    s.append(Paragraph("Filed under A.R.S. 19-121, 19-121.01 and 19-121.04. Record held by the Elections "
                       "Division, Arizona Secretary of State, 1700 West Washington Street, Phoenix.",
                       st["small"]))
    d = _doc(path, "Filing and sample selection record: initiative serial %s" % K.SERIAL, K.OFFICE)
    d.build(s, onFirstPage=_footer("%s %s" % (K.OFFICE, K.DIVISION)),
            onLaterPages=_footer("%s %s" % (K.OFFICE, K.DIVISION)))
    PP.normalize_pdf(path, b"D:20280711084200-07'00'", b"D:20280711163500-07'00'")


# ---------------------------------------------------------------------------------- the instruction
def instruction_text():
    return [
        ("What is transmitted to you", [
            "You will receive a copy of every sampled signature line assigned to your county, with the "
            "registration number of the circulator who gathered the sheet, the number printed on that "
            "sheet and the line number on it. Those three together identify the line, because a sheet "
            "number is only that circulator's own series and every circulator's runs from one.",
        ]),
        ("Giving a line its disposition", [
            "Verify each line against your county's registration records as they stood on the date written "
            "on the line, which is the date the elector signed. Give every line exactly one disposition code "
            "from the table below. Work down the table and record the first code that applies.",
            "A signature shall not be counted more than once. Where an elector's signature appears on more "
            "than one of the lines before you, count the earliest by the date signed and give each later "
            "line code DU.",
        ]),
        ("Returning your determinations", [
            "Return one row for each line transmitted to you, carrying the line's identification, the "
            "disposition code, the registration identifier you matched the line to where you matched one, "
            "and the date of your determination. You may return in more than one transmittal; mark each with "
            "your return reference.",
            "The Elections Division runs a consolidation check over the returned lines. Where that check "
            "finds a line that should have carried code DU, or where a line is returned without a "
            "disposition, the division records the determination and reports the line as determined by the "
            "division rather than by the recorder.",
            "The division may send you a query notice asking you to re-determine a named line. A "
            "re-determination replaces the disposition you first returned for that line and nothing else.",
        ]),
    ]


def write_instruction(path):
    st = styles()
    s = [Paragraph("Signature verification instruction for county recorders", st["h1"]),
         Paragraph("%s, %s. Issued %s. Applies to every statewide initiative, referendum and recall "
                   "petition sample transmitted on or after 1 January 2028. Supersedes the instruction of "
                   "9 March 2024." % (K.OFFICE, K.DIVISION, K.long_date(K.INSTRUCTION_DATE)), st["sub"])]
    for head, paras in instruction_text():
        s.append(Paragraph(esc(head), st["h2"]))
        for pr in paras:
            s.append(Paragraph(esc(pr), st["body"]))
    s.append(Spacer(1, 4))
    s.append(_table(st, ["Code", "Disposition"], [[c, K.CODE_TEXT[c]] for c in K.CODE_KEYS],
                    [0.7 * inch, 4.4 * inch]))
    s.append(Paragraph("Issued by %s, State Elections Director. Questions to the Elections Division, "
                       "Arizona Secretary of State." % K.PEOPLE["director"], st["small"]))
    d = _doc(path, "Signature verification instruction for county recorders", K.OFFICE)
    d.build(s, onFirstPage=_footer("%s %s" % (K.OFFICE, K.DIVISION)),
            onLaterPages=_footer("%s %s" % (K.OFFICE, K.DIVISION)))
    PP.normalize_pdf(path, b"D:20271114103100-07'00'", b"D:20271118161500-07'00'")


# ---------------------------------------------------------------------------------- Navajo's own code list
def oldcodes_text():
    out = [
        "NAVAJO COUNTY RECORDER",
        "Petition signature verification: disposition codes",
        "Issued %s by %s, Navajo County Recorder. Holbrook, Arizona." % (K.us_date(K.OLD_CODES_DATE),
                                                                         K.PEOPLE["recorder"]),
        "",
        "Enter one code per line in the DISPOSITION column of the return sheet.",
        "",
    ]
    for code, text in K.OLD_CODES:
        out.append("  %s   %s" % (code, text))
    out += [
        "",
        "Where a line is entered under code 02, record in the REGISTRATION column the registration",
        "identifier the line was matched to if one was matched, and leave that column empty if none",
        "was.",
        "",
        "This list is the recorder's own. It has not been reissued since %d." % K.OLD_CODES_DATE.year,
    ]
    return out


def write_oldcodes(path):
    with open(path, "w", encoding="utf-8", newline="\r\n") as fh:
        fh.write("\n".join(oldcodes_text()) + "\n")


# ---------------------------------------------------------------------------------- the 2024 close-out
def closeout_county_rows(w):
    """The examination's own administration, county by county: the lines each team took, the examiner days
    it spent on them and the day it closed. The close-out is the record of the examination, not a table of
    validity rates."""
    r = K.rng("closeout", seed=w.seed)
    shares = [s for _, _, s in K.COUNTIES]
    tot = sum(shares)
    out, left = [], K.PRIOR_LINES
    rate = K.PRIOR_LINES / float(K.PRIOR_DAYS * K.PRIOR_STAFF)
    for i, (cc, name, share) in enumerate(K.COUNTIES):
        if i == len(K.COUNTIES) - 1:
            ex = left
        else:
            ex = int(round(K.PRIOR_LINES * share / tot * (1 + 0.06 * (r.random() - 0.5))))
            ex = min(ex, left - (len(K.COUNTIES) - i - 1) * 60)
        left -= ex
        edays = K.round_half_up(ex / rate * (1 + 0.09 * (r.random() - 0.5)), 1)
        closed = K.PRIOR_CLOSEOUT - dt.timedelta(days=int(r.integers(0, 17)))
        out.append([cc, name, ex, edays, K.iso(closed)])
    return out


def write_closeout(path, w):
    import openpyxl
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    rows = closeout_county_rows(w)
    r2 = K.rng("closeout_total", seed=w.seed)
    notc = int(round(K.PRIOR_LINES * (0.274 + 0.004 * (r2.random() - 0.5))))
    counted = K.PRIOR_LINES - notc
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Full count"
    head = [
        ("Measure serial", K.PRIOR_MEASURE),
        ("Measure", "Arizona Rural Broadband Access Act"),
        ("Petition filed", K.iso(dt.date(2024, 7, 3))),
        ("Signature lines filed", K.PRIOR_LINES),
        ("Full count ordered", K.iso(K.PRIOR_ORDERED)),
        ("Order", "Maricopa County Superior Court, cause CV2024-009718"),
        ("Examination", "Every signature line filed was examined; no sample was drawn"),
        ("Examination opened", K.iso(K.PRIOR_OPENED)),
        ("Examination closed", K.iso(K.PRIOR_CLOSEOUT)),
        ("Examination business days", K.PRIOR_DAYS),
        ("Examiners assigned", K.PRIOR_STAFF),
        ("Signature lines examined", K.PRIOR_LINES),
        ("Signature lines counted", counted),
        ("Signature lines not counted", notc),
        ("Signatures required", K.PRIOR_REQUIRED),
        ("Determination", "Insufficient"),
        ("Certified", K.iso(K.PRIOR_CERTIFIED)),
        ("Close-out prepared by", K.PEOPLE["analyst"]),
    ]
    for k, v in head:
        ws.append([k, v])
    ws.append([])
    ws.append(["County", "County name", "Lines examined", "Examiner days", "Examination closed"])
    for r in rows:
        ws.append(r)
    for i, wd in enumerate([30, 24, 16, 15, 20], start=1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    for c in ws[1]:
        c.font = Font(bold=True)
    for c in ws[len(head) + 2]:
        c.font = Font(bold=True)
        c.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(path)
    PP.normalize_ooxml(path, K.PEOPLE["analyst"], K.PEOPLE["analyst"], "2024-08-29T15:41:00Z",
                       "2024-08-29T16:12:00Z", title="%s full count close-out" % K.PRIOR_MEASURE,
                       excel=True, company=K.OFFICE)
    return dict(counted=counted, not_counted=notc, rows=rows)


def write_all(out, w, n):
    write_record(os.path.join(out, F_RECORD))
    write_instruction(os.path.join(out, F_INSTRUCTION))
    write_oldcodes(os.path.join(out, F_OLDCODES))
    return write_closeout(os.path.join(out, R.F_CLOSEOUT), w)
