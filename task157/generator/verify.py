"""task157 (architecture 2) independent verifier: reads only the shipped bundle, on its own code path with
pandas, and recomputes the projection, the statutory outcome, every rung of the ladder, every ask answer and
the full count's own sizing, then compares them with generator/figures.json and the two goldens. It imports
no generator module, and it takes the statutory figures out of the filing record's own text rather than from
a constant typed here.

    python3 task157/generator/verify.py <target>"""
from __future__ import annotations

import datetime as dt
import json
import math
import os
import re
import sys

import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
FAIL = []
RUN = []

RETURNS = "county_verification_returns_2028.csv"
CORR = "county_return_corrections_2028.csv"
SNAP_LATE = "avid_snapshot_extract_20280710.csv"
SNAP_EARLY = "avid_snapshot_extract_20271220.csv"
SHEETS = "petition_sheet_index_2028.csv"
CIRC = "circulator_registration_extract_i18-2028.csv"
MANIFEST = "circulator_sheet_submissions_2028.csv"
RECORD = "sos_filing_and_selection_record_i18-2028.pdf"
INSTRUCTION = "county_recorder_verification_instruction_2027.pdf"
OLDCODES = "navajo_recorder_disposition_codes_2019.txt"
CLOSEOUT = "measure_i2024-04_full_count_closeout.xlsx"

SINGLE = ["NR", "CA", "AD", "SG", "NC", "NQ", "IL", "MD", "WS"]
CODES = ["V"] + SINGLE[:6] + ["DU"] + SINGLE[6:]
OLD_ONE_TO_ONE = {"01": "V", "03": "AD", "04": "SG", "05": "NC", "06": "NQ", "07": "DU", "08": "IL",
                  "09": "MD", "10": "WS"}
DMY_COUNTY, OLD_COUNTY = "09", "10"
MONTHS = ["2028-%02d" % m for m in range(1, 7)]
MONS = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]


def ok(label, cond, detail=""):
    print("%s  %-70s %s" % ("PASS" if cond else "FAIL", label, detail))
    RUN.append(label)
    if not cond:
        FAIL.append(label)


def nearest(x, step):
    return int(step * math.floor(x / step + 0.5))


def pdf_text(path):
    import pypdf
    return re.sub(r"\s+", " ", "\n".join(p.extract_text() or "" for p in pypdf.PdfReader(path).pages))


def parse_date(county, text):
    if text is None or (isinstance(text, float) and text != text):
        return None
    t = str(text).strip()
    if not t or t == "nan":
        return None
    if county == DMY_COUNTY:
        d, m, y = t.split("/")
        return dt.date(int(y), int(m), int(d))
    if county == OLD_COUNTY:
        d, mon, y = t.split("-")
        return dt.date(2000 + int(y), MONS.index(mon.upper()) + 1, int(d))
    return dt.date(*(int(x) for x in t.split("-")))


def main(target):
    p = lambda f: os.path.join(target, f)
    figs = json.load(open(os.path.join(HERE, "figures.json")))

    # ---------------------------------------------------------------- the statute, from the record itself
    rec = pdf_text(p(RECORD))
    ok("the filing record reads back through this verifier's own extractor", len(rec) > 2000,
       "%d characters" % len(rec))
    filed = int(re.search(r"receipted 38,080 sheets, bearing ([\d,]+) signature lines",
                          rec).group(1).replace(",", ""))
    sample = int(re.search(r"random sample of five per cent of the signature lines filed, ([\d,]+) of the",
                           rec).group(1).replace(",", ""))
    required = int(re.search(r"the number required for serial I-18-2028 is therefore ([\d,]+)",
                             rec).group(1).replace(",", ""))
    lo = int(re.search(r"ninety-five per cent of the number required, ([\d,]+) signatures",
                       rec).group(1).replace(",", ""))
    hi = int(re.search(r"one hundred and five per cent of the number required, ([\d,]+) signatures",
                       rec).group(1).replace(",", ""))
    ok("the statutory figures come out of the record's own text",
       (filed, sample, required, lo, hi) == (figs["lines_filed"], figs["sample_lines"], figs["required"],
                                            figs["band_lo"], figs["band_hi"]),
       "filed %d, sample %d, required %d, band %d to %d" % (filed, sample, required, lo, hi))
    factor = filed // sample
    ok("the sample's own factor divides the filed total exactly",
       factor * sample == filed and factor == figs["sample_factor"], "%d" % factor)
    instr = pdf_text(p(INSTRUCTION))
    ok("the verification instruction reads back and carries the order of the grounds",
       "Work down the table and record the first code that applies" in instr, "quoted")
    oldtxt = open(p(OLDCODES), encoding="utf-8").read()
    ok("Navajo's own code list reads back and carries its merged code",
       "not registered or registration cancelled" in oldtxt, "code 02")

    # ---------------------------------------------------------------- conform the returns
    r = pd.read_csv(p(RETURNS), dtype={"county_code": str, "voter_reg_id": "Int64"})
    ok("the returns carry one row per sampled signature line", len(r) == sample, "%d rows" % len(r))
    late = pd.read_csv(p(SNAP_LATE), dtype={"county_code": str, "voter_reg_id": "Int64"})
    cor = pd.read_csv(p(CORR), dtype={"county_code": str})
    key = ["circulator_reg_no", "sheet_no", "line_no"]
    r = r.merge(cor[key + ["corrected_disposition_code"]], on=key, how="left")
    r["code"] = r["corrected_disposition_code"].fillna(r["disposition_code"])
    navajo = r["county_code"].eq(OLD_COUNTY)
    r.loc[navajo, "code"] = r.loc[navajo, "code"].map(
        lambda c: OLD_ONE_TO_ONE.get(c, c) if c != "02" else "02")
    split = navajo & r["code"].eq("02")
    r.loc[split, "code"] = r.loc[split, "voter_reg_id"].notna().map({True: "CA", False: "NR"})
    ok("the corrections supersede the disposition first returned, and nothing else",
       int(r["corrected_disposition_code"].notna().sum()) == len(cor),
       "%d query notices applied" % len(cor))
    ok("every sampled line carries exactly one code from the current set",
       set(r["code"].unique()) <= set(CODES) and len(r) == sample,
       "%d codes over %d lines" % (r["code"].nunique(), len(r)))

    single = int(r["code"].isin(SINGLE).sum())
    dup = int(r["code"].eq("DU").sum())
    ok("the sampled counts match the write-up",
       (single, dup) == (figs["sampled_single"], figs["sampled_dup"]),
       "%d on other grounds, %d duplicates" % (single, dup))

    # ---------------------------------------------------------------- the projection
    single_proj = single * factor
    dup_proj = dup * factor * factor
    valid = filed - single_proj - dup_proj
    band = "insufficient" if valid < lo else ("sufficient" if valid > hi else "full count")
    ok("the single signature grounds carry at the sample's own factor",
       single_proj == figs["projected_single"], "%d lines" % single_proj)
    ok("the duplicate strikes carry at the square of the sampling fraction",
       dup_proj == figs["projected_dup"] and factor * factor == figs["pair_denominator"],
       "%d lines, one pair in %d" % (dup_proj, factor * factor))
    ok("the projection and its printed cell match the write-up",
       valid == figs["call_valid"] and nearest(valid, figs["bin"]) == figs["call_printed"],
       "%d prints %d" % (valid, nearest(valid, figs["bin"])))
    ok("the projection as a share of the number required matches the write-up",
       abs(round(100.0 * valid / required, 1) - figs["call_pct"]) < 1e-9,
       "%.1f per cent" % (100.0 * valid / required))
    ok("the statutory outcome is the examination of all signature lines filed", band == "full count", band)

    # ---------------------------------------------------------------- the ladder
    early = pd.read_csv(p(SNAP_EARLY), dtype={"voter_reg_id": "Int64"})
    act = set(early.loc[early["registration_status"].eq("ACTIVE"), "voter_reg_id"].dropna().astype(int))
    slip = int(r.loc[r["code"].eq("CA"), "voter_reg_id"].dropna().astype(int).isin(act).sum())
    rung0 = filed - (single - slip) * factor - dup * factor
    rung1 = filed - single * factor - dup * factor
    county_dup = int((r["code"].eq("DU") & r["determined_by"].eq("RECORDER")).sum())
    rung2 = filed - single * factor - (county_dup * factor + (dup - county_dup) * factor * factor)
    ok("rung 0, the default path, reproduces and files a different band",
       rung0 == figs["default_path_valid"] and nearest(rung0, figs["bin"]) == figs["default_path_printed"]
       and rung0 > hi, "%d prints %d, sufficient" % (rung0, nearest(rung0, figs["bin"])))
    ok("rung 1 reproduces and files a different band",
       rung1 == figs["rung1_valid"] and rung1 > hi, "%d prints %d" % (rung1, nearest(rung1, figs["bin"])))
    ok("rung 2 reproduces", rung2 == figs["rung2_valid"], "%d prints %d"
       % (rung2, nearest(rung2, figs["bin"])))
    ok("the decisive step is the largest step in the ladder",
       (rung2 - valid) > (rung0 - rung1) and (rung2 - valid) > (rung1 - rung2),
       "%d against %d and %d" % (rung2 - valid, rung0 - rung1, rung1 - rung2))

    # ---------------------------------------------------------------- the registration state
    late = late.set_index("voter_reg_id")
    reg_date = late["registration_date"].to_dict()
    canc_date = late["cancellation_date"].to_dict()
    status = late["registration_status"].to_dict()

    def iso(x):
        if x is None or x != x or str(x).strip() == "":
            return None
        return dt.date(*(int(v) for v in str(x).split("-")))

    r["signed"] = [parse_date(c, t) for c, t in zip(r["county_code"], r["signature_date"])]

    def in_force(vid, when):
        if when is None or vid is None:
            return False
        rd, cd = iso(reg_date.get(vid)), iso(canc_date.get(vid))
        if rd is None or rd > when:
            return False
        if cd is None or rd > cd:
            return True
        return when < cd

    r["lapsed"] = [bool(pd.notna(v) and in_force(int(v), d) and status.get(int(v)) == "CANCELLED")
                   for v, d in zip(r["voter_reg_id"], r["signed"])]

    # ---------------------------------------------------------------- ask A, per county
    det = r[r["determined_by"].eq("RECORDER")]
    bad = []
    for cc, g in det.groupby("county_code"):
        if int(len(g)) != figs["ask_a_%s_determined" % cc] or int(g["lapsed"].sum()) != \
                figs["ask_a_%s_lapsed" % cc]:
            bad.append(cc)
    ok("every county's two figures reproduce", not bad, "15 counties" if not bad else ",".join(bad))
    ok("the lines the division determined are outside every county's figures",
       int(len(det)) + int((~r["determined_by"].eq("RECORDER")).sum()) == sample,
       "%d determined by the division" % int((~r["determined_by"].eq("RECORDER")).sum()))

    # ---------------------------------------------------------------- ask B, per firm and month
    cr = pd.read_csv(p(CIRC))
    cr["start"] = cr["registered_on"].map(iso)
    cr["end"] = cr["withdrawn_on"].map(iso)
    by_circ = {}
    for row in cr.itertuples():
        by_circ.setdefault(row.circulator_reg_no, []).append((row.start, row.end,
                                                              row.firm_registration_no))
    for v in by_circ.values():
        v.sort(key=lambda t: t[0])

    def firm_on(circ, when):
        for start, end, firm in by_circ.get(circ, []):
            if start > when or (end is not None and end <= when):
                continue
            return firm
        return None

    inwin = r[r["signed"].map(lambda d: d is not None and dt.date(2028, 1, 3) <= d <= dt.date(2028, 6, 30))]
    cells = {}
    for row in inwin.itertuples():
        f = firm_on(row.circulator_reg_no, row.signed)
        if f is None:
            continue
        k = (f, "%04d-%02d" % (row.signed.year, row.signed.month))
        c0 = cells.setdefault(k, [0, 0])
        c0[0] += 1
        if row.code in SINGLE:
            c0[1] += 1
    bad = []
    for k in sorted({k for k in figs if k.startswith("ask_b_")}):
        m = re.match(r"ask_b_(C-\d{4}-\d{4})_(\d{4}-\d{2})_(signatures|struck)", k)
        got = cells.get((m.group(1), m.group(2)), [0, 0])[0 if m.group(3) == "signatures" else 1]
        if got != figs[k]:
            bad.append(k)
    ok("every firm and month cell reproduces on the as-of registration", not bad,
       "%d cells" % (len(figs) and len([k for k in figs if k.startswith("ask_b_")]) // 2)
       if not bad else ",".join(bad[:4]))
    ok("the circulators who registered again under another firm are in the extract",
       int((cr["registration_status"].eq("WITHDRAWN")).sum()) == figs["circulator_moves"],
       "%d withdrawn rows" % int((cr["registration_status"].eq("WITHDRAWN")).sum()))

    # ---------------------------------------------------------------- ask C, per month
    sh = pd.read_csv(p(SHEETS))
    ser = {(a, b): c for a, b, c in zip(sh["circulator_reg_no"], sh["sheet_no"], sh["sheet_serial"])}
    bad = []
    for m in MONTHS:
        sel = r[r["signed"].map(lambda d, m=m: d is not None and "%04d-%02d" % (d.year, d.month) == m)]
        serials = {ser[(a, b)] for a, b in zip(sel["circulator_reg_no"], sel["sheet_no"])}
        if int(len(sel)) != figs["ask_c_%s_signatures" % m] or len(serials) != figs["ask_c_%s_sheets" % m]:
            bad.append(m)
    ok("every month's two figures reproduce on the office's own sheet serials", not bad,
       "6 months" if not bad else ",".join(bad))
    ok("a sheet number alone does not identify a sheet",
       sh["sheet_no"].nunique() < len(sh) and len(ser) == len(sh),
       "%d sheet numbers over %d sheets" % (sh["sheet_no"].nunique(), len(sh)))

    # ---------------------------------------------------------------- the filed total's other reading
    man = pd.read_csv(p(MANIFEST))
    claimed = int(man["signatures_claimed"].sum())
    alt = claimed - single * (claimed / float(sample)) - dup / ((sample / float(claimed)) ** 2)
    ok("the manifests' own total is ruled out by the record and files the same cell",
       "struck at receipt; they are not part of the filing" in rec
       and nearest(alt, figs["bin"]) == figs["call_printed"],
       "%d claimed, projecting %d" % (claimed, alt))

    # ---------------------------------------------------------------- the full count's own sizing
    import openpyxl
    wb = openpyxl.load_workbook(p(CLOSEOUT), data_only=True)
    ws = wb["Full count"]
    head = {}
    for row in ws.iter_rows(values_only=True):
        if row and isinstance(row[0], str) and len(row) >= 2:
            head[row[0]] = row[1]
    rate = head["Signature lines examined"] / float(head["Examination business days"]
                                                    * head["Examiners assigned"])
    days = int(math.ceil((filed / rate) / 30.0))
    ok("the 2024 close-out reads back and sizes the full count inside the deadline",
       head["Determination"] == "Insufficient" and days <= 20,
       "%.1f lines an examiner-day, %d business days at 30 examiners" % (rate, days))
    wb.close()

    # ---------------------------------------------------------------- the goldens
    gb = os.path.join(TASK, "golden", "signature_projection_2028.xlsx")
    wb = openpyxl.load_workbook(gb, data_only=True)
    vals = []
    for name in wb.sheetnames:
        for row in wb[name].iter_rows(values_only=True):
            vals.extend([v for v in row if v is not None])
    wb.close()
    ok("the golden workbook carries the projection, the printed cell and the outcome",
       figs["call_valid"] in vals and figs["call_printed"] in vals
       and "Examine all signature lines filed" in vals, "the determination sheet")
    ok("the golden workbook carries every ask figure",
       all(figs[k] in vals for k in figs if k.startswith("ask_")),
       "%d ask figures" % len([k for k in figs if k.startswith("ask_")]))
    png = os.path.join(TASK, "golden", "petition_disposition_2028.png")
    ok("the golden chart is a png with no writing library in its bytes",
       open(png, "rb").read(8) == b"\x89PNG\r\n\x1a\n"
       and b"matplotlib" not in open(png, "rb").read().lower(), "%d bytes" % os.path.getsize(png))

    print()
    print("verify: %d checks, %d failed" % (len(RUN), len(FAIL)))
    for f in FAIL:
        print("  FAILED: %s" % f)
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(TASK, "target")))
