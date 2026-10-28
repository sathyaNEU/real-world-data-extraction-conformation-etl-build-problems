"""task162 independent verifier. Reads only the bundle, on its own code path (pandas over the CSVs, pypdf over the PDFs,
every filed number parsed out of the documents rather than imported), and recomputes every graded figure: the Phase 3
call, the twelve Phase 4 candidates' 2025-26 leakage and the zone's 2025-26 balance. It also reads back every format the
pack ships, each with a positive control, and compares everything with generator/figures.json.

    python3 task162/generator/verify.py <target>"""
from __future__ import annotations

import datetime as dt
import email
import json
import math
import os
import re
import sys

import numpy as np
import openpyxl
import pandas as pd
import pypdf

HERE = os.path.dirname(os.path.abspath(__file__))
FAILS = []


def check(label, cond, detail=""):
    print("%s  %s%s" % ("PASS" if cond else "FAIL", label, ("  [%s]" % (detail,)) if detail != "" else ""))
    if not cond:
        FAILS.append(label)


def pdf(path):
    return "\n".join(p.extract_text() for p in pypdf.PdfReader(path).pages)


def one(text):
    return " ".join(text.split())


def num(pattern, text):
    m = re.search(pattern, text)
    if not m:
        raise SystemExit("filed figure not found: %s" % pattern)
    return float(m.group(1))


def rnd(x, nd):
    q = x * 10 ** nd
    return "%.*f" % (nd, math.floor(q + 0.5) / 10 ** nd)


def main(target):
    t = lambda f: os.path.join(target, f)
    fig = json.load(open(os.path.join(HERE, "figures.json"), encoding="utf-8"))

    # ---------------------------------------------------------------- documents: every format read back, controls
    raw_m = pdf(t("LW-LA-04_DMA_leakage_assessment_method.pdf"))
    meth, sched = one(raw_m), one(pdf(t("outcome_contract_schedule_4_phase3.pdf")))
    cert = one(pdf(t("TH-OUT-0417_calibration_certificate.pdf")))
    with open(t("2027-28_planning_assumptions.eml"), "rb") as fh:
        msg = email.message_from_binary_file(fh)
    note = one(msg.get_payload(decode=True).decode("utf-8"))
    vs = openpyxl.load_workbook(t("phase3_prv_schedule_rev_C.xlsx"), read_only=True).active
    vrows = list(vs.iter_rows(values_only=True))
    ex = openpyxl.load_workbook(t("dma_leakage_assessment_phase3_2025_26.xlsx"), read_only=True)
    check("V1 pdf: LW-LA-04 reads back (control: the valid-day sentence)", "A day counts when the district's inlet" in meth)
    check("V1 pdf: Schedule 4 reads back (control: section 4.4)", "find-and-fix gang-year" in sched)
    check("V1 pdf: the certificate reads back (control: the as-found line)", "As found: reading" in cert)
    check("V1 eml: the planning note reads back (control: the sender and the rev C sentence)",
          "Clive Heath" in msg["From"] and "settings in Annette's Phase 3 schedule (rev C)" in note)
    check("V1 xlsx: the valve schedule and the export read back (controls: the header rows)",
          any(r and r[0] == "District" for r in vrows) and ex.sheetnames == ["Monthly", "Year"])

    # filed numbers, parsed from the documents
    target_red = num(r"target reduction is (\d\.\d+) Ml/d", sched)
    gang = num(r"at (\d\.\d+) Ml/d of sustained reduction for each find-and-fix gang-year", sched)
    hh_allow = num(r"allowed at (\d\.\d+) litres an hour for each dwelling", meth)
    nhh_allow = num(r"and (\d\.\d+) litres an hour for each non-household", meth)
    night = int(num(r"The night hour is (\d\d):00 to", meth))
    bar = num(r"1 bar is (\d+\.\d+) metres head", meth)
    unb = num(r"Unbilled authorised use is (\d\.\d+) per cent", meth) / 100
    ureg = num(r"under-registration at (\d\.\d+) per cent", meth) / 100
    unauth = num(r"unauthorised use at (\d\.\d+) per cent", meth) / 100
    err = num(r"reading (\d\.\d+) per cent low", cert) / 100
    cal = dt.datetime.strptime(re.search(r"Calibration date: (\d+ \w+ \d{4})", cert).group(1), "%d %B %Y").date()
    lines = [x.strip() for x in raw_m[raw_m.find("Table C2"):].splitlines()]
    n1 = {lines[i]: float(lines[i + 2]) for i in range(len(lines) - 2)
          if re.fullmatch(r"NZ\d{2}", lines[i]) and re.fullmatch(r"\d\.\d{2}", lines[i + 2])}
    s_txt = pdf(t("outcome_contract_schedule_4_phase3.pdf"))
    phase3 = sorted(set(re.findall(r"NZ\d{2}", s_txt[s_txt.find("Table 4.1"):])))
    hi = [i for i, r in enumerate(vrows) if r and r[0] == "District"][0]
    col = vrows[hi].index("Outlet setting (m)")
    setting = {r[0]: float(r[col]) for r in vrows[hi + 1:] if r and r[0]}
    check("V2 the filed figures parse: target, gang-year, allowances, night hour, shares, certificate, N1, settings",
          len(n1) == 34 and len(phase3) == 22 and set(setting) == set(phase3),
          "target %.2f, gang %.2f, night %02d:00, N1 %d, Phase 3 %d" % (target_red, gang, night, len(n1), len(phase3)))

    # ---------------------------------------------------------------- telemetry into hourly frames
    reg = pd.read_csv(t("logger_register.csv"), dtype=str)
    codes = pd.read_csv(t("district_code_history.csv"), dtype=str).set_index("old_code")["current_code"].to_dict()
    dreg = pd.read_csv(t("district_register.csv"), dtype=str)
    inlets_n = dreg.set_index("district_code")["inlets"].astype(int).to_dict()
    reg["district"] = reg["district_code"].map(lambda c: codes.get(c, c))
    reg["own"] = reg["district"] == reg["district_code"]
    t0 = pd.Timestamp("2025-04-01")

    def load(f):
        d = pd.read_csv(t(f), dtype={"logger_serial": str}).drop_duplicates(["logger_serial", "reading_time_gmt"])
        d = d.merge(reg[["logger_serial", "district", "own", "pressure_unit"]], on="logger_serial", how="left")
        d["p_m"] = np.where(d["pressure_unit"] == "bar", d["pressure"] * bar, d["pressure"])
        k = ((pd.to_datetime(d["reading_time_gmt"]) - t0) / pd.Timedelta(hours=1)).astype(int)
        d["day"], d["hour"] = k // 24, k % 24
        return d

    inl, zon = load("northern_zone_inlet_loggers_2025_26.csv"), load("northern_zone_pressure_loggers_2025_26.csv")
    ndays = 365
    full = pd.MultiIndex.from_product([range(ndays), range(24)], names=["day", "hour"])
    flow, pin, azp = {}, {}, {}
    for c, g in inl.groupby("district"):
        agg = g.groupby(["day", "hour"]).agg(q=("flow_l_s", "sum"), n=("flow_l_s", "size"))
        q = agg["q"].where(agg["n"] == inlets_n[c]).reindex(full)
        flow[c] = q.to_numpy().reshape(ndays, 24)
        own = g[g["own"]].set_index(["day", "hour"])["p_m"].reindex(full)
        pin[c] = own.to_numpy().reshape(ndays, 24)
    for c, g in zon.groupby("district"):
        azp[c] = g.set_index(["day", "hour"])["p_m"].reindex(full).to_numpy().reshape(ndays, 24)
    # a district fed through a neighbour carries no inflow all day; it and the neighbour sharing the open boundary
    # valve are not closed that day
    bv = {r["district_code"]: set(v.strip() for v in r["boundary_valves"].split(";")) for _, r in dreg.iterrows()}
    not_closed = {c: set() for c in flow}
    for c, q in flow.items():
        dead = np.where(np.isfinite(q).all(axis=1) & (np.nanmax(np.where(np.isnan(q), 1e9, q), axis=1) < 0.1))[0]
        for d in dead:
            not_closed[c].add(int(d))
            for o in flow:
                if o != c and bv[o] & bv[c]:
                    not_closed[o].add(int(d))
    check("V3 the open boundary is found in the flows: NZ09 dry and NZ10 not closed on the same nine days",
          sorted(not_closed["NZ09"]) == sorted(not_closed["NZ10"]) and len(not_closed["NZ09"]) == 9
          and not [c for c in flow if not_closed[c] and c not in ("NZ09", "NZ10")])

    # ---------------------------------------------------------------- dwellings and non-households by day
    prop = pd.read_csv(t([f for f in os.listdir(target) if f.startswith("property_extract")][0]), dtype=str).fillna("")
    prop["district"] = prop["district_code"].map(lambda c: codes.get(c, c))

    def dayno(s, default):
        return ((pd.to_datetime(s, errors="coerce") - pd.Timestamp("2025-04-01")).dt.days).fillna(default)

    prop["a"] = dayno(prop["connected"], -1).astype(int)
    prop["b"] = dayno(prop["disconnected"], 10 ** 6).astype(int)
    span = np.arange(ndays)
    dw, nh = {}, {}
    for c, g in prop.groupby("district"):
        hh = g[g["property_class"] == "household"]
        on = lambda x: ((span[None, :] >= x["a"].to_numpy()[:, None]) & (span[None, :] < x["b"].to_numpy()[:, None]))
        tot = np.zeros(ndays)
        for _, sp in hh.groupby("supply_point_ref"):
            with_u = on(sp[sp["uprn"] != ""]).sum(0)
            without = on(sp[sp["uprn"] == ""]).sum(0)
            tot += with_u + np.where(with_u > 0, 0, without)
        dw[c] = tot
        nh[c] = on(g[g["property_class"] == "non-household"]).sum(0)

    # ---------------------------------------------------------------- LW-LA-04 day by day
    months = [dt.date(2025 + (3 + i) // 12, (3 + i) % 12 + 1, 1) for i in range(12)]
    mstart = [(d - dt.date(2025, 4, 1)).days for d in months] + [ndays]

    def year(per_day, valid):
        ms = [per_day[mstart[i]:mstart[i + 1]][valid[mstart[i]:mstart[i + 1]]].mean() for i in range(12)]
        return float(np.mean(ms))

    def terms(c):
        q, a = flow[c], azp[c]
        valid = ~np.isnan(q).any(1) & ~np.isnan(a).any(1) & ~np.isnan(pin[c]).any(1)
        valid[list(not_closed[c])] = False
        nnf = q[:, night] - (dw[c] * hh_allow + nh[c] * nhh_allow) / 3600.0
        return valid, nnf, a, a[:, night]

    def leak(nnf, a, aznp, n):
        with np.errstate(invalid="ignore"):
            return nnf * 3.6 * np.sum((a / aznp[:, None]) ** n, axis=1) / 1000.0

    # the Phase 3 call
    yr = {r[0]: float(r[2]) for r in ex["Year"].iter_rows(min_row=4, values_only=True) if r and r[0] and str(r[0]).startswith("NZ")}
    base = sum(yr[c] for c in phase3)
    fc = 0.0
    for c in phase3:
        valid, nnf, a, aznp = terms(c)
        out = np.minimum(pin[c], setting[c])
        a2 = a - (pin[c] - out)
        with np.errstate(invalid="ignore"):
            nnf2 = nnf * (a2[:, night] / aznp) ** n1[c]
        fc += year(leak(nnf2, a2, a2[:, night], n1[c]), valid)
    red = base - fc
    short = max(0.0, target_red - red)
    gy = int(math.ceil(short / gang - 1e-9)) if short > 0 else 0
    F = fig["call"]
    check("V4 the Phase 3 call recomputes: 2025-26, forecast, reduction, shortfall, gang-years",
          (rnd(base, 2), rnd(fc, 2), rnd(red, 2), rnd(short, 2), gy) ==
          (F["base"], F["forecast"], F["reduction"], F["shortfall"], F["gang_years"]),
          "%.4f %.4f %.4f %.4f %d" % (base, fc, red, short, gy))
    check("V4b the 2025-26 baseline recomputes from the loggers as the export states it",
          abs(sum(year(leak(*terms(c)[1:], n1[c]), terms(c)[0]) for c in phase3) - base) < 0.0005)

    # Ask A
    near = sorted(set(dreg["district_code"]) - set(phase3))
    got = {}
    for c in near:
        valid, nnf, a, aznp = terms(c)
        got[c] = rnd(year(leak(nnf, a, aznp, n1[c]), valid), 3)
    check("V5 the twelve Phase 4 candidates recompute", got == fig["answers"]["ask_a"],
          {c: (got[c], fig["answers"]["ask_a"][c]) for c in near if got[c] != fig["answers"]["ask_a"][c]})

    # Ask B
    ou = pd.read_csv(t("thornholme_outlet_daily_2025_26.csv"))
    ou["d"] = pd.to_datetime(ou["date"]).dt.date
    ou = ou[(ou["d"] >= dt.date(2025, 4, 1)) & (ou["d"] <= dt.date(2026, 3, 31))]
    siv = float(np.where(ou["d"] < cal, ou["outlet_volume_ml"] / (1 - err), ou["outlet_volume_ml"]).sum()) / ndays
    zone = set(dreg["district_code"]) | {"NZ-TM"}
    sps = set(prop.loc[prop["district"].isin(zone), "supply_point_ref"])
    bl = pd.read_csv(t("billing_extract_north_2025_26.csv"), dtype={"supply_point_ref": str})
    bl = bl[bl["supply_point_ref"].isin(sps)].copy()
    f0, f1 = pd.to_datetime(bl["bill_from"]), pd.to_datetime(bl["bill_to"])
    y0, y1 = pd.Timestamp("2025-04-01"), pd.Timestamp("2026-03-31")
    inside = ((np.minimum(f1, y1) - np.maximum(f0, y0)).dt.days + 1).clip(lower=0)
    bl["v"] = bl["volume_m3"] * inside / ((f1 - f0).dt.days + 1)
    met = bl.loc[bl["basis"] == "metered", "v"].sum() / 1000 / ndays
    unm = bl.loc[bl["basis"] != "metered", "v"].sum() / 1000 / ndays
    real = siv - met - unm - unb * siv - (ureg * met + unauth * siv)
    B = fig["answers"]["ask_b"]
    check("V6 the zone balance recomputes: system input, billed unmetered, real losses (and billed metered)",
          (rnd(siv, 2), rnd(unm, 2), rnd(real, 2), rnd(met, 2)) ==
          (B["system_input"], B["billed_unmetered"], B["real_losses"], B["billed_metered"]),
          "%.4f %.4f %.4f %.4f" % (siv, unm, real, met))
    check("V7 csv: every shipped CSV parses with a header and rows",
          all(len(pd.read_csv(t(f), nrows=5).columns) >= 2 for f in os.listdir(target) if f.endswith(".csv")))
    print("\nverify: %d failed" % len(FAILS))
    return 1 if FAILS else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
