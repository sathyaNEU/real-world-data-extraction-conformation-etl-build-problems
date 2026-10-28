"""task162 checks on the built pack: the call and its corridor, the three rules executed on the generator's clean truth,
every semantic fork row, the ladder, both asks with their devices and rungs, the separation line, the shipped
sentences, the realism sweep, the input gates and the distractors, and the goldens against figures.json. Run by build.py
after the goldens are written; exits 1 on any failure.

    python3 task162/generator/checks.py"""
from __future__ import annotations

import builtins
import copy
import csv
import datetime as dt
import email
import json
import os
import re
import subprocess
import sys
import zipfile

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import calc as C            # noqa: E402
import common as K          # noqa: E402
import deliver as DL        # noqa: E402
import docs as DOC          # noqa: E402
import records as R         # noqa: E402
import world as W           # noqa: E402

TASK = os.path.dirname(HERE)
TARGET = os.path.join(TASK, "target")
GOLDEN = os.path.join(TASK, "golden")
RESULTS = []


def ok(label, cond, detail=""):
    RESULTS.append((bool(cond), label))
    print("%s  %s%s" % ("PASS" if cond else "FAIL", label, ("  [%s]" % (detail,)) if detail != "" else ""))


def p2(x):
    return "%.2f" % K.round_half_up(x, 2)


def p3(x):
    return "%.3f" % K.round_half_up(x, 3)


def flat(s):
    return " ".join(s.split())


# ------------------------------------------------------------------------------------------------- readers
def pdf_text(path):
    return subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True, check=True).stdout


def xlsx_text(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            out.append(" ".join(str(v) for v in row if v is not None))
    wb.close()
    return "\n".join(out)


def eml_text(path):
    with open(path, "rb") as fh:
        m = email.message_from_binary_file(fh)
    return m.get_payload(decode=True).decode("utf-8")


def texts():
    out = {}
    for f in sorted(os.listdir(TARGET)):
        p = os.path.join(TARGET, f)
        if f.endswith(".pdf"):
            out[f] = flat(pdf_text(p))
        elif f.endswith(".xlsx"):
            out[f] = flat(xlsx_text(p))
        elif f.endswith(".eml"):
            out[f] = flat(eml_text(p))
    return out


def rows(f):
    with open(os.path.join(TARGET, f), encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


# ------------------------------------------------------------------------------------------------- the clean truth
class Clean(C.Pack):
    """The generator's truth before any record is written: every hour present, every pressure in metres, one inflow per
    district (NZ04's two inlets one stream), each district's dwellings and non-households as the world counts them."""

    def __init__(self, T, N1, S):
        self.N1, self.SETTINGS, self.t0 = N1, S, K.YEAR_START
        self.flow, self.azp, self.pin, self.hh, self.nhh = {}, {}, {}, {}, {}
        for p in T["P"]:
            D = T[p["code"]]
            self.flow[p["code"]] = D["flow"].reshape(C.N_DAYS, 24)
            self.azp[p["code"]] = D["azp"].reshape(C.N_DAYS, 24)
            self.pin[p["code"]] = D["inlet"].reshape(C.N_DAYS, 24)
            self.hh[p["code"]] = D["hh_day"].astype(float)
            self.nhh[p["code"]] = np.full(C.N_DAYS, float(p["nonhh"]))

    def households_by_day(self, code, mode="dwellings", map_codes=True, cls="household"):
        return self.hh[code] if cls == "household" else self.nhh[code]


def call_of(P, codes=W.PHASE3, mode="replay", base=None):
    fc = sum(float(P.forecast(c, mode).mean()) for c in codes)
    b = base if base is not None else sum(float(P.assess(c)[0].mean()) for c in codes)
    red = b - fc
    short = max(0.0, K.TARGET - red)
    return dict(base=b, fc=fc, reduction=red, shortfall=short, met=red >= K.TARGET,
                gangs=K.ceil_div(short, K.GANG_YIELD) if short > 0 else 0)


def printed(c):
    return (p2(c["fc"]), p2(c["reduction"]), p2(c["shortfall"]), c["gangs"])


# ------------------------------------------------------------------------------------------------- main
def main():
    figs = json.load(open(os.path.join(HERE, "figures.json"), encoding="utf-8"))
    meta = json.load(open(os.path.join(TASK, "metadata.json"), encoding="utf-8"))
    TX = texts()
    files = sorted(os.listdir(TARGET))

    # ---------------------------------------------------------------- positive controls (every reader reads)
    controls = {DOC.F_METHOD: "The night hour is 03:00 to 04:00 GMT, the hourly mean stamped 03:00.",
                DOC.F_SCHEDULE: "The Phase 3 target reduction is 0.48 Ml/d.",
                DOC.F_CERT: "As found: reading 1.4 per cent low across the flow range.",
                DOC.F_NOTE: "with no growth.",
                DOC.F_VALVES: "Settings are metres head at the valve outlet, the point the district's inlet logger reads.",
                R.F_EXPORT: "Lynmere Water: DMA leakage assessment, 2025-26, Phase 3 districts (LW-LA-04 issue 6)"}
    for f, s in controls.items():
        ok("P1 %s reads back and carries its known sentence" % f, len(TX.get(f, "")) > 200 and s in TX[f])

    # ---------------------------------------------------------------- input gates and distractors
    big = max((sum(1 for _ in open(os.path.join(TARGET, f), encoding="utf-8")) - 1) for f in files if f.endswith(".csv"))
    ok("G1 ten or more files", len(files) >= 10, len(files))
    ok("G2 three or more formats", len({f.rsplit(".", 1)[1] for f in files}) >= 3, sorted({f.rsplit(".", 1)[1] for f in files}))
    ok("G3 a file of 25,000 rows or more", big >= 25000, big)
    dis = meta["distractor_files"]
    ok("G4 two or more distractors, each shipped", len(dis) >= 2 and all(d in files for d in dis), dis)
    bad = [f for f in files if re.search(r"distractor|readme|dictionary|notes|manifest|schema|known_issues", f, re.I)]
    ok("G5 no file named for its role or describing the data", not bad, bad)
    import openpyxl
    sheets = [s for f in files if f.endswith(".xlsx") for s in openpyxl.load_workbook(os.path.join(TARGET, f), read_only=True).sheetnames]
    ok("G6 no sheet describes the data", not [s for s in sheets if re.search(r"note|readme|about|dictionar|field|legend", s, re.I)], sheets)

    opened = set()
    real_open = builtins.open

    def spy(path, *a, **k):
        opened.add(os.path.basename(str(path)))
        return real_open(path, *a, **k)
    builtins.open = spy
    try:
        X = DL.compute(TARGET)
    finally:
        builtins.open = real_open
    ok("G7 the golden and both asks read neither distractor (deleting either changes nothing)",
       not (set(dis) & opened) and R.F_BILLS in opened and R.F_INLET in opened, sorted(opened & set(files)))

    # ---------------------------------------------------------------- the call
    N1, S = C.read_n1(TARGET), C.read_settings(TARGET)
    P = C.Pack(TARGET, N1, S)
    T = W.simulate()
    ok("C0 N1 and the settings read off the shipped PDF and workbook match the world",
       N1 == {p["code"]: p["n1"] for p in T["P"]} and S == {p["code"]: p["setting"] for p in T["P"] if p["phase"] == 3})
    fc = figs["call"]
    ok("C1 the call from the pack is figures.json's (2025-26, forecast, reduction, shortfall, gang-years, not met)",
       (p2(X["base"]), p2(X["fc"]), p2(X["reduction"]), p2(X["shortfall"]), X["gangs"]) ==
       (fc["base"], fc["forecast"], fc["reduction"], fc["shortfall"], fc["gang_years"]) and fc["meets_target"] == "no"
       and X["reduction"] < K.TARGET,
       "%s %s %s %s %d" % (fc["base"], fc["forecast"], fc["reduction"], fc["shortfall"], fc["gang_years"]))
    base_x = X["base"]
    rec = call_of(P)
    ok("C2 SF11 the export's 2025-26 baseline and the method rerun on the loggers print alike",
       p2(rec["base"]) == p2(base_x) and abs(rec["base"] - base_x) < 0.0005, "%.5f %.5f" % (rec["base"], base_x))
    clear = {k: K.bin_clearance(X[k], 0.01) for k in ("base", "fc", "reduction", "shortfall")}
    ok("C3 every printed call figure clears its cell edge by a quarter step or more",
       min(clear.values()) >= 0.25, {k: round(v, 2) for k, v in clear.items()})
    ok("C3b the gang-years clear their whole-number boundary by a quarter or more",
       K.ceil_clearance(X["shortfall"], K.GANG_YIELD) >= 0.25, round(K.ceil_clearance(X["shortfall"], K.GANG_YIELD), 2))
    gold_p = (p2(X["fc"]), p2(X["reduction"]), p2(X["shortfall"]), X["gangs"])
    for mode, sf in (("feedback", "SF12"), ("open_loss", "SF10")):
        c = call_of(P, mode=mode, base=base_x)
        ok("C4 %s the %s corridor prints the golden's figures" % (sf, mode), printed(c) == gold_p, printed(c))
    dflt = call_of(P, mode="nightratio", base=base_x)
    ok("C5 position: the default path sits 0.10 Ml/d or more over the target and the golden 0.06 or more under it",
       dflt["reduction"] >= K.TARGET + 0.10 and X["reduction"] <= K.TARGET - 0.06,
       "%.4f %.4f" % (dflt["reduction"], X["reduction"]))
    ok("C6 the default path's reduction is figures.json's rival figure", p2(dflt["reduction"]) == fc["default_reduction"],
       p2(dflt["reduction"]))

    # ---------------------------------------------------------------- the ladder
    rung = {"R0 no change": dict(reduction=0.0, gangs=K.ceil_div(K.TARGET, K.GANG_YIELD), met=False)}
    for name, mode in (("R1 night ratio (default)", "nightratio"), ("R1b export and schedule AZNP", "export_ratio"),
                       ("R2a night flow times 24", "night24"), ("R2b night drop every hour", "shift"),
                       ("R3 typical day replayed", "typday")):
        rung[name] = call_of(P, mode=mode, base=base_x)
    for name in ("R1 night ratio (default)", "R1b export and schedule AZNP", "R2a night flow times 24",
                 "R2b night drop every hour"):
        ok("L %s files the target met, no gang-years" % name, rung[name]["met"] and rung[name]["gangs"] == 0,
           "%.4f" % rung[name]["reduction"])
    ok("L R3 typical day replayed misses the golden's printed reduction (the averaged day, against the per-day factor)",
       p2(rung["R3 typical day replayed"]["reduction"]) != p2(X["reduction"]),
       "%.4f, %d gang-years" % (rung["R3 typical day replayed"]["reduction"], rung["R3 typical day replayed"]["gangs"]))
    ok("L R0 no change files %d gang-years" % rung["R0 no change"]["gangs"], rung["R0 no change"]["gangs"] == 10)
    ok("L separation: no rung prints the golden's reduction, and every rung below R3 files another call",
       all(p2(v["reduction"]) != p2(X["reduction"]) for v in rung.values())
       and all(v["gangs"] != X["gangs"] for k, v in rung.items() if not k.startswith("R3")),
       {k.split()[0]: (p2(v["reduction"]), v["gangs"]) for k, v in rung.items()})
    saved = dict(P.N1)
    for c in W.PHASE3:
        P.N1[c] = 1.15
    one = call_of(P, base=base_x)
    P.N1.update(saved)
    c2 = re.findall(r"(NZ\d{2})\s+\S.*?\s+(\d\.\d{2})", TX[DOC.F_METHOD])
    ok("L SF6 one company N1 (1.15) moves the call off 2 gang-years; Table C2 lists every district's own",
       one["gangs"] != X["gangs"] and len({a for a, _ in c2}) == 34, "%d gang-years, %d rows" % (one["gangs"], len({a for a, _ in c2})))

    # ---------------------------------------------------------------- the three rules on the clean truth
    Q = Clean(T, N1, S)
    qg = call_of(Q)
    qd = call_of(Q, mode="nightratio", base=qg["base"])
    ok("R1a the field's raw material ships: hourly inflow and pressures at every Phase 3 inlet and average zone point",
       all(np.isfinite(P.flow[c]).sum() > 8000 and np.isfinite(P.azp[c]).sum() > 8000 and np.isfinite(P.pin[c]).sum() > 8000
           for c in W.PHASE3))
    import openpyxl as _ox
    vcols = [c.value for c in list(_ox.load_workbook(os.path.join(TARGET, DOC.F_VALVES), read_only=True).active.iter_rows(min_row=5, max_row=5))[0]]
    ecols = [c.value for c in list(_ox.load_workbook(os.path.join(TARGET, R.F_EXPORT), read_only=True)["Monthly"].iter_rows(min_row=4, max_row=4))[0]]
    ok("R1b no shipped column holds the deciding quantity: the valve schedule and the export carry no 2027-28 or saving "
       "column", vcols == DOC.VALVE_COLS and not [c for c in vcols + ecols if re.search(r"2027|saving|benefit|forecast", str(c), re.I)],
       vcols)
    ok("R1c the export's monthly columns alone cannot reach the golden: the night ratio on them files the target met",
       rung["R1b export and schedule AZNP"]["met"] and abs(rung["R1b export and schedule AZNP"]["reduction"] - X["reduction"]) > 0.15)
    ok("R3a the analyst test on the clean truth: the night ratio with every filed rule files the target met",
       qd["met"] and qd["gangs"] == 0, "%.4f" % qd["reduction"])
    ok("R3b the decisive step alone (each hour through the setting) files the golden's call on the clean truth",
       not qg["met"] and qg["gangs"] == X["gangs"], "%.4f %d" % (qg["reduction"], qg["gangs"]))
    for mode in ("shift", "night24", "typday"):
        q = call_of(Q, mode=mode, base=qg["base"])
        ok("R3b a lower rung's single correction (%s) stays wrong on the clean truth" % mode,
           p2(q["reduction"]) != p2(qg["reduction"]) and (mode == "typday" or q["gangs"] != X["gangs"]),
           "%.4f %d" % (q["reduction"], q["gangs"]))

    # R2: every Phase 3 mess item, mishandled the natural way, leaves the call and its printed figures
    def mutate(fn):
        Pm = copy.copy(P)
        Pm.flow, Pm.azp, Pm.pin = ({c: a.copy() for c, a in d.items()} for d in (P.flow, P.azp, P.pin))
        fn(Pm)
        return Pm

    def interp(a):
        f = a.reshape(-1).copy()
        i, g = np.arange(f.size), np.isfinite(f)
        f[~g] = np.interp(i[~g], i[g], f[g])
        return f.reshape(a.shape)

    bst = np.array([K.bst(K.YEAR_START + dt.timedelta(days=d, hours=12)) for d in range(C.N_DAYS)])

    def local_night(Pm):
        for c in W.PHASE3:
            for d in (Pm.flow, Pm.azp, Pm.pin):
                f = d[c].reshape(-1)
                d[c] = np.where(bst[:, None], np.concatenate([[np.nan], f[:-1]]).reshape(C.N_DAYS, 24), d[c])

    swap = min(r["removed"] for r in rows(R.F_LOGGERS) if r["district_code"] == "NZ20" and r["removed"])

    def nz20_old_serial(Pm):
        k = (dt.date.fromisoformat(swap) - K.YEAR_START.date()).days
        Pm.flow["NZ20"][:k + 1] = np.nan
        Pm.pin["NZ20"][:k + 1] = np.nan

    def gaps_interp(Pm):
        Pm.azp["NZ27"] = interp(Pm.azp["NZ27"])
        Pm.flow["NZ15"], Pm.pin["NZ15"] = interp(Pm.flow["NZ15"]), interp(Pm.pin["NZ15"])

    def resends_summed(Pm):
        k = (dt.date(2025, 6, 17) - K.YEAR_START.date()).days
        Pm.flow["NZ22"][k:k + 2] *= 2.0

    def swap_interp(Pm):
        Pm.flow["NZ20"], Pm.pin["NZ20"] = interp(Pm.flow["NZ20"]), interp(Pm.pin["NZ20"])

    for name, fn in (("loggers' GMT stamps read as local time (SF4)", local_night),
                     ("NZ20's old inlet serial left unmapped", nz20_old_serial),
                     ("battery gaps at NZ27 and NZ15 interpolated and kept (SF8)", gaps_interp),
                     ("NZ22's re-sent inlet rows summed with the originals", resends_summed),
                     ("NZ20's swap-day hours interpolated and kept", swap_interp)):
        Pm = mutate(fn)
        fcm = sum(float(np.nanmean(Pm.forecast(c))) for c in W.PHASE3)
        red = base_x - fcm
        sh = max(0.0, K.TARGET - red)
        got = (p2(fcm), p2(red), p2(sh), K.ceil_div(sh, K.GANG_YIELD))
        ok("R2 %s leaves the call and its printed figures" % name, got == gold_p, got)
    Pb = copy.copy(P)
    Pb.azp = {c: (a / C.BAR if c in ("NZ03", "NZ06", "NZ11") else a) for c, a in P.azp.items()}
    ok("R2 the bar loggers (NZ06 and NZ11 zone points, NZ03 after its swap) left unconverted leave Ask A",
       all(p3(v) == figs["answers"]["ask_a"][c] for c, v in C.ask_a(Pb, ["NZ06", "NZ11"]).items()))
    ok("R2 the bar make reads nowhere in Phase 3 (every Phase 3 logger reports metres head)",
       all(r["pressure_unit"] == "metres head" for r in rows(R.F_LOGGERS) if r["district_code"] in W.PHASE3))

    # R3c: no plant
    blocks = 0
    for f, col in ((R.F_INLET, "pressure"), (R.F_ZONE, "pressure")):
        ser = {}
        for r in rows(f):
            ser.setdefault(r["logger_serial"], {})[r["reading_time_gmt"]] = r[col]
        for s_, d in ser.items():
            v = [d[k] for k in sorted(d)]
            seen = set()
            for i in range(0, len(v) - 24, 24):
                key = tuple(v[i:i + 24])
                blocks += key in seen
                seen.add(key)
    ok("R3c no logger series repeats a day's block of readings at another time (no copied or backfilled block)",
       blocks == 0, blocks)
    docs_text = " ".join(TX[f] for f in (DOC.F_METHOD, DOC.F_SCHEDULE, DOC.F_NOTE, DOC.F_VALVES, DOC.F_CERT))
    ok("R3c no document states a 2027-28 leakage, saving or forecast figure",
       not re.search(r"2027-28 (leakage|reduction)[^.]{0,60}\d\.\d+ Ml/d|forecast (leakage|reduction)|saving of|benefit of",
                     docs_text, re.I))
    az_after = {r[0]: float(r[6]) for r in _ox.load_workbook(os.path.join(TARGET, DOC.F_VALVES), read_only=True).active.iter_rows(min_row=6, values_only=True) if r[0]}
    mine = {}
    for c in W.PHASE3:
        okd = P.valid(c)
        a_n, p_n = P.azp[c][okd, K.NIGHT_HOUR], P.pin[c][okd, K.NIGHT_HOUR]
        mine[c] = float(np.mean(a_n - (p_n - np.minimum(p_n, S[c]))))
    ok("R3c the valve schedule's AZNP after commissioning is a true night figure, not a planted one",
       all(abs(az_after[c] - mine[c]) < 0.051 for c in W.PHASE3))

    # ---------------------------------------------------------------- the mechanism as the note states it
    dayg = [c for c in W.PHASE3 if int(c[2:]) % 2 == 1 or int(c[2:]) in (14, 22, 30)]
    share, drop, depth, wk, tot = {}, [], [], 0, 0
    for c in W.PHASE3:
        okd = P.valid(c)
        op = (P.pin[c] <= S[c]) & okd[:, None]
        share[c] = float(op.sum()) / (24.0 * okd.sum())
        wkd = np.array([(K.YEAR_START + dt.timedelta(days=d)).weekday() < 5 for d in range(C.N_DAYS)])
        drop.append(float(np.nanmean(P.pin[c][okd, K.NIGHT_HOUR]) - np.nanmean(P.pin[c][okd & wkd, 9:16])))
        depth.append(float(np.nanmean(P.pin[c][okd, K.NIGHT_HOUR]) - S[c]))
        for d, h in zip(*np.where(op)):
            lt = K.local(K.YEAR_START + dt.timedelta(days=int(d), hours=int(h)))
            wk += lt.weekday() < 5 and 6 <= lt.hour < 22
            tot += 1
    sd = [share[c] for c in dayg]
    sp = [share[c] for c in W.PHASE3 if c not in dayg]
    ok("D1 the valves would stand open 20 to 37 per cent of the hours in 14 districts and 3 to 15 in the other 8",
       len(sd) == 14 and 0.195 <= min(sd) and max(sd) < 0.375 and 0.025 <= min(sp) and max(sp) < 0.155,
       "%.3f to %.3f; %.3f to %.3f" % (min(sd), max(sd), min(sp), max(sp)))
    ok("D2 the open hours are weekday hours from 06:00 to 22:00 local (the golden's stated mechanism), every valve has them",
       tot and wk / tot >= 0.99 and min(share.values()) > 0, "%.3f of %d hours" % (wk / tot, tot))
    ok("D3 the weekday daytime inlet sits 7 to 11 m under the night inlet, the settings 8 to 15 m under it",
       6.5 <= min(drop) and max(drop) < 11.5 and 7.5 <= min(depth) and max(depth) < 15.5,
       "%.1f to %.1f; %.1f to %.1f" % (min(drop), max(drop), min(depth), max(depth)))

    # ---------------------------------------------------------------- semantic fork rows
    s4 = TX[DOC.F_SCHEDULE]
    ok("SF1 the reduction is in leakage as LW-LA-04 assesses it (Schedule 4.2), not night flow",
       "each assessed by method LW-LA-04 issue 6" in s4 and rung["R2a night flow times 24"]["met"])
    ok("SF2 the forward conditions are filed: each month as recorded in the same month of 2025-26, no growth, valves all year",
       "as they were recorded in the same month of 2025-26" in TX[DOC.F_NOTE] and "with no growth" in TX[DOC.F_NOTE]
       and "for the whole of 2027-28" in TX[DOC.F_NOTE])
    days = {}
    for c in W.PHASE3:
        okd = P.valid(c)
        t = P.day_terms(c, P.households_by_day(c))
        days[c] = (t["leak"][okd].mean(), P.forecast_days(c)[okd].mean())
    alt = (sum(v[0] for v in days.values()), sum(v[1] for v in days.values()))
    red3 = alt[0] - alt[1]
    sh3 = max(0, K.TARGET - red3)
    ok("SF3 the year as the mean of all valid days rather than of the twelve months: same call",
       K.ceil_div(sh3, K.GANG_YIELD) == X["gangs"] and not red3 >= K.TARGET and "The year's leakage is the mean of its twelve months." in TX[DOC.F_METHOD],
       "%.4f" % red3)
    ok("SF4 the night hour is filed in GMT (section 3) and a local reading prints the same (R2)",
       "The night hour is 03:00 to 04:00 GMT" in TX[DOC.F_METHOD])
    agree = all(all(np.array_equal(P.households_by_day(c), P.households_by_day(c, mode=m)) for m in ("rows", "supply_points", "uprn"))
                for c in W.PHASE3)
    ok("SF5 Phase 3's dwellings count the same however the property extract is counted", agree)
    Pz = copy.copy(P)

    def zone_point(code, mode="replay", S_=None):
        s_ = S[code]
        n1 = P.N1[code]
        okd = P.valid(code)
        t = P.day_terms(code, P.households_by_day(code))
        a2 = np.minimum(P.azp[code], s_)
        aznp2 = a2[:, K.NIGHT_HOUR]
        with np.errstate(invalid="ignore"):
            nnf2 = t["nnf"] * (aznp2 / t["aznp"]) ** n1
            ndf2 = np.sum((a2 / aznp2[:, None]) ** n1, axis=1)
        return P.monthly(nnf2 * 3.6 * ndf2 / 1000.0, okd)
    Pz.forecast = zone_point
    zp = call_of(Pz, base=base_x)
    ok("SF7 the setting read at the average zone point instead of the outlet prints the same call; the schedule files "
       "the outlet", printed(zp) == gold_p and controls[DOC.F_VALVES] in TX[DOC.F_VALVES], printed(zp))
    part = 0.0
    for c in W.PHASE3:
        f, a = P.flow[c], P.azp[c]
        okd = ~np.isnan(f[:, K.NIGHT_HOUR]) & ~np.isnan(a[:, K.NIGHT_HOUR]) & ~np.isnan(P.pin[c][:, K.NIGHT_HOUR])
        a2 = a - (P.pin[c] - np.minimum(P.pin[c], S[c]))
        hh, nhh = P.households_by_day(c), P.households_by_day(c, cls="non-household")
        nnf = f[:, K.NIGHT_HOUR] - (hh * K.ALLOWANCE + nhh * K.NONHH_ALLOWANCE) / 3600.0
        with np.errstate(invalid="ignore"):
            nnf2 = nnf * (a2[:, K.NIGHT_HOUR] / a[:, K.NIGHT_HOUR]) ** P.N1[c]
            ndf2 = np.nansum((a2 / a2[:, [K.NIGHT_HOUR]]) ** P.N1[c], axis=1)
        part += float(np.nanmean(P.monthly(np.where(okd, nnf2 * 3.6 * ndf2 / 1000.0, np.nan), okd)))
    red8 = base_x - part
    ok("SF8 partial days kept (missing hours left out of the factor): same call",
       K.ceil_div(max(0, K.TARGET - red8), K.GANG_YIELD) == X["gangs"], "%.4f" % red8)
    mnf_fc = 0.0
    for c in W.PHASE3:
        okd = P.valid(c)
        t = P.day_terms(c, P.households_by_day(c))
        a2 = P.azp[c] - (P.pin[c] - np.minimum(P.pin[c], S[c]))
        aznp2 = a2[:, K.NIGHT_HOUR]
        with np.errstate(invalid="ignore"):
            nnf2 = t["mnf"] * (aznp2 / t["aznp"]) ** P.N1[c] - t["allow"]
            mnf_fc += float(P.monthly(nnf2 * 3.6 * np.sum((a2 / aznp2[:, None]) ** P.N1[c], axis=1) / 1000.0, okd).mean())
    ok("SF13 N1 applied to the whole night flow, night use included, files the target met; section 7 applies it to "
       "leakage and section 4 nets the night use first",
       base_x - mnf_fc >= K.TARGET and "to any change in a district's leakage with its pressure" in TX[DOC.F_METHOD]
       and "The net night flow is the minimum night flow less the allowance" in TX[DOC.F_METHOD], "%.4f" % (base_x - mnf_fc))
    near = int(K.round_half_up(X["shortfall"] / K.GANG_YIELD))
    ok("SF9 gang-years rounded to the nearest would file %d; Schedule 4 files them rounded up" % near,
       "rounded up to whole gang-years" in s4 and K.ceil_div(X["shortfall"], K.GANG_YIELD) == X["gangs"], near)

    # ---------------------------------------------------------------- Ask A
    A = figs["answers"]["ask_a"]
    gold_a = X["ask_a"]
    ok("A1 the twelve candidates' figures are figures.json's", {c: p3(v) for c, v in gold_a.items()} == A)
    ok("A1b every candidate's figure clears its cell edge by 0.3 of a step or more",
       min(K.bin_clearance(v, 0.001) for v in gold_a.values()) >= 0.3, round(min(K.bin_clearance(v, 0.001) for v in gold_a.values()), 2))
    bins = []
    for v_, need in (("rows", 12), ("supply_points", 12), ("uprn", 12)):
        got = C.ask_a(P, W.NEAR, v_)
        miss = sum(p3(got[c]) != A[c] for c in W.NEAR)
        d = [abs(round((float(p3(got[c])) - float(A[c])) * 1000)) for c in W.NEAR]
        bins += d
        ok("A2 dwellings counted by %s miss %d of 12" % (v_, need), miss >= need, "%d missed, %d to %d bins" % (miss, min(d), max(d)))
    ok("A2b every wrong count misses each figure by 1 to 7 bins (the note's range)", min(bins) >= 1 and max(bins) <= 7,
       "%d to %d" % (min(bins), max(bins)))
    got = C.ask_a(P, W.NEAR, "unmapped")
    ok("A3 NZ04A's properties left unmapped miss NZ04 and nothing else",
       [c for c in W.NEAR if p3(got[c]) != A[c]] == ["NZ04"])
    got = C.ask_a(P, W.NEAR, "open_kept")
    ok("A4 NZ09's and NZ10's open-boundary days kept miss both and nothing else",
       [c for c in W.NEAR if p3(got[c]) != A[c]] == ["NZ09", "NZ10"])
    P2 = C.Pack(TARGET, N1, S, skip_codes=("NZ04A",))
    ok("A5 NZ04's second inlet (logged as NZ04A) left out misses NZ04", p3(C.ask_a(P2, ["NZ04"])["NZ04"]) != A["NZ04"])
    props = rows(R.F_PROPS)
    by_sp = {}
    for p in props:
        by_sp.setdefault(p["supply_point_ref"], []).append(p)
    plots = [p for p in props if p["property_ref"].startswith("PL-")]
    paired = closed = 0
    lags = []
    for p in plots:
        sib = [q for q in by_sp[p["supply_point_ref"]] if q is not p]
        if len(sib) == 1 and sib[0]["uprn"] and sib[0]["connected"]:
            paired += 1
            lag = (dt.date.fromisoformat(sib[0]["connected"]) - dt.date.fromisoformat(p["connected"])).days
            lags.append(lag)
            closed += p["disconnected"] == sib[0]["connected"]
            closed -= bool(p["disconnected"]) and p["disconnected"] != sib[0]["connected"]
    ok("A6 every new build's plot record shares its supply point with one address record 45 to 150 days later",
       paired == len(plots) and lags and min(lags) >= 45 and max(lags) <= 150, "%d plots, lag %d to %d" % (len(plots), min(lags), max(lags)))
    unclosed = sum(1 for p in plots if not p["disconnected"]) / float(len(plots))
    ok("A7 a closed plot record closes on its address record's connection date; about one in six is closed",
       closed == sum(1 for p in plots if p["disconnected"]) and 0.8 <= unclosed <= 0.87, round(unclosed, 3))
    shared = [v for v in by_sp.values() if len([q for q in v if q["uprn"]]) >= 2 and not any(q["property_ref"].startswith("PL-") for q in v)]
    ok("A8 flats on a shared supply: distinct UPRNs, households, assessed, connected throughout",
       shared and all(len({q["uprn"] for q in v}) == len(v) and all(q["property_class"] == "household" and q["charging_basis"] == "unmeasured"
                                                                    and not q["connected"] for q in v) for v in shared), len(shared))
    codes_map = {r["old_code"]: r["current_code"] for r in rows(R.F_CODES)}
    p3props = [p for p in props if codes_map.get(p["district_code"], p["district_code"]) in W.PHASE3]
    ok("A9 separation line: no plot record, shared supply or retired code in Phase 3 (its infill and demolitions are "
       "single records every count reads alike, SF5)",
       not [p for p in p3props if p["property_ref"].startswith("PL-") or not p["uprn"]
            or p["district_code"] not in W.PHASE3 or len(by_sp[p["supply_point_ref"]]) > 1]
       and any(p["connected"] for p in p3props) and any(p["disconnected"] for p in p3props), len(p3props))
    k0 = (dt.date(2026, 1, 12) - K.YEAR_START.date()).days
    low = np.nanmax(P.flow["NZ09"][:, :], axis=1) < 0.2
    ok("A10 NZ09's inlet carries nothing on exactly the nine days 12 to 20 January 2026 and on no other day",
       list(np.where(low)[0]) == list(range(k0, k0 + 9)), list(np.where(low)[0]))
    ok("A11 separation line: no open boundary, unit change or retired-code logger in Phase 3",
       not [r for r in rows(R.F_LOGGERS) if r["district_code"] in W.PHASE3 and r["pressure_unit"] != "metres head"]
       and not [r for r in rows(R.F_CODES) if r["current_code"] in W.PHASE3 and r["effective"] >= "2025-04-01"])

    # ---------------------------------------------------------------- Ask B
    B = figs["answers"]["ask_b"]
    g = X["ask_b"]
    ok("B1 the balance is figures.json's (system input, billed unmetered, real losses and the rest)",
       (p2(g["siv"]), p2(g["unmeasured"]), p2(g["real"]), p2(g["metered"]), p2(g["unbilled"]), p2(g["apparent"])) ==
       (B["system_input"], B["billed_unmetered"], B["real_losses"], B["billed_metered"], B["unbilled"], B["apparent"]))
    keys = ("siv", "unmeasured", "real", "metered", "unbilled", "apparent")
    for v_ in ("multiply", "exclusive"):
        h = C.ask_b(TARGET, v_)
        ok("B2 the %s corridor prints every balance figure alike" % v_, all(p2(h[k]) == p2(g[k]) for k in keys),
           {k: p2(h[k]) for k in keys})
    lo = C.ask_b(TARGET, "multiply")
    ex = C.ask_b(TARGET, "exclusive")
    both = g["real"] + (lo["real"] - g["real"]) + (ex["real"] - g["real"])
    ok("B2b both corridor readings at once print real losses alike", p2(both) == p2(g["real"]), "%.4f" % both)
    for v_, hit in (("dedupe", ("unmeasured", "real")), ("no_certificate", ("siv", "real")),
                    ("codes_unmapped", ("unmeasured", "real")), ("no_trunk", ("real",)), ("bill_dates", ("real",))):
        h = C.ask_b(TARGET, v_)
        kept = [k for k in ("siv", "unmeasured", "real") if k not in hit]
        ok("B3 %s misses %s and keeps the rest" % (v_, " and ".join(hit)),
           all(p2(h[k]) != p2(g[k]) for k in hit) and all(p2(h[k]) == p2(g[k]) for k in kept),
           {k: p2(h[k]) for k in ("siv", "unmeasured", "real")})
    bills = rows(R.F_BILLS)
    dup = {}
    for b in bills:
        if b["basis"] == "assessed":
            dup[tuple(b.values())] = dup.get(tuple(b.values()), 0) + 1
    cnt = {}
    for p in props:
        if p["charging_basis"] == "unmeasured":
            cnt[p["supply_point_ref"]] = cnt.get(p["supply_point_ref"], 0) + 1
    multi = {k: n for k, n in dup.items() if n > 1}
    ok("B4 every set of identical assessed lines sits on a supply point with that many assessed dwellings or more",
       multi and all(cnt.get(k[0], 0) >= n for k, n in multi.items()), "%d sets" % len(multi))
    ok("B5 section 8 files the balance's every share and the customers (NZ-TM named)",
       all(s_ in TX[DOC.F_METHOD] for s_ in ("Unbilled authorised use is 0.4 per cent of system input",
                                             "under-registration at 2.5 per cent of billed metered consumption",
                                             "each bill apportioned to the year by days", "(code NZ-TM in the property system)")))

    # ---------------------------------------------------------------- shipped sentences: pins once, no announcement
    pins = {"0.48 Ml/d": DOC.F_SCHEDULE, "0.05 Ml/d": DOC.F_SCHEDULE, "1.7 litres an hour": DOC.F_METHOD,
            "03:00 to 04:00 GMT": DOC.F_METHOD, "with no growth": DOC.F_NOTE, "10.197": DOC.F_METHOD,
            "Settings are metres head at the valve outlet": DOC.F_VALVES}
    for s_, where in pins.items():
        n = {f: TX[f].count(s_) for f in TX}
        ok("S1 the pin '%s' appears once, in %s" % (s_, where), n[where] == 1 and sum(n.values()) == 1, {f: v for f, v in n.items() if v})
    words = r"regulat|stands? (fully )?open|fully open|wide open|below (its|the|their) setting|lower of|hour by hour|replay|\bbite|bites|throttl"
    hits = {f: re.findall(words, t, re.I) for f, t in TX.items() if f != R.F_EXPORT}
    ok("S2 no shipped document says when a valve acts or how a scheme's effect is forecast",
       not any(hits.values()) and len(docs_text) > 3000, {f: v for f, v in hits.items() if v})
    ok("S2b the search above reads the documents it searches (positive control on the same extractor)",
       re.search(r"Fixed-outlet pressure reducing valves", TX[DOC.F_VALVES]) is not None)

    # ---------------------------------------------------------------- realism sweep
    def arithmetic(seq):
        d = np.diff(np.array(seq, dtype=float))
        return len(d) > 3 and np.all(d == d[0])
    nums = {"property_ref": [int(p["property_ref"][1:]) for p in props if not p["property_ref"].startswith("PL-")][:2000],
            "uprn": [int(p["uprn"]) for p in props if p["uprn"]][:2000],
            "supply_point_ref": [int(p["supply_point_ref"][2:]) for p in props][:2000],
            "contact_ref": [int(r["contact_ref"][3:]) for r in rows(R.F_CONTACTS)],
            "logger_serial": [int(r["logger_serial"].split("-")[1]) for r in rows(R.F_LOGGERS)]}
    ok("Z1 no identifier column runs at a constant step", not [k for k, v in nums.items() if arithmetic(v)])
    dreg = rows(R.F_DISTRICTS)
    ok("Z2 district codes do not follow the names' alphabetical order",
       [r["district_name"] for r in dreg] != sorted(r["district_name"] for r in dreg))
    import openpyxl as ox
    mrows = list(ox.load_workbook(os.path.join(TARGET, R.F_EXPORT), read_only=True)["Monthly"].iter_rows(min_row=5, values_only=True))
    vd = [r[2] for r in mrows if r and r[0]]
    ok("Z3 valid-day counts vary across districts and months", len(set(vd)) >= 5, sorted(set(vd))[:8])
    ok("Z4 no round headline total", abs(X["base"] * 100 - round(X["base"] * 100)) > 1e-6 and abs(g["siv"] * 10 - round(g["siv"] * 10)) > 1e-3)
    people = set(K.PEOPLE.values())
    named = set(re.findall(r"\b(Emily Fitzgerald|Annette Patel|Clive Heath|Louis Lee|Mandy Austin|[A-Z][a-z]+ [A-Z][a-z]+(?= <))", " ".join(TX.values())))
    ok("Z5 every person in the pack is a drawn persona", named and named <= people, sorted(named))
    leaks = []
    for f in files:
        p = os.path.join(TARGET, f)
        if f.endswith(".xlsx"):
            with zipfile.ZipFile(p) as z:
                for n in z.namelist():
                    if re.search(rb"openpyxl|python", z.read(n), re.I):
                        leaks.append(f + ":" + n)
        elif f.endswith(".pdf"):
            if re.search(rb"ReportLab|reportlab|AAAAAA\+", open(p, "rb").read()):
                leaks.append(f)
    ok("Z6 containers carry no writing library's signature", not leaks, leaks)
    late = [p for p in props if max(p["connected"] or "0", p["disconnected"] or "0") > K.AS_OF.isoformat()]
    late += [b for b in bills if b["bill_to"] > K.AS_OF.isoformat()]
    ok("Z7 no extract carries a date after its 2 November 2026 pull", not late, len(late))
    tel_last = max(r["reading_time_gmt"] for r in rows(R.F_ZONE))
    ok("Z7b the telemetry and the contact log end with the assessment year", tel_last <= "2026-03-31 23:00"
       and max(r["received"] for r in rows(R.F_CONTACTS)) <= "2026-03-31 23:59", tel_last)
    pdf_dates = {f: re.search(rb"/CreationDate \(D:(\d{8})", open(os.path.join(TARGET, f), "rb").read()) for f in files if f.endswith(".pdf")}
    ok("Z8 no document is dated after the files were taken", all(m and m.group(1).decode() <= K.AS_OF.strftime("%Y%m%d") for m in pdf_dates.values())
       and "19 Oct 2026" in open(os.path.join(TARGET, DOC.F_NOTE), encoding="utf-8").read())
    inl = {}
    for r in rows(R.F_INLET):
        d = r["reading_time_gmt"][:10]
        inl[d] = inl.get(d, 0.0) + float(r["flow_l_s"]) * 3.6 / 1000.0
    mill = {}
    for r in rows(R.F_MILL):
        d = r["reading_time_gmt"][:10]
        mill[d] = mill.get(d, 0.0) + float(r["flow_l_s"]) * 3.6 / 1000.0
    outl = {r["date"]: float(r["outlet_volume_ml"]) / (1.0 + K.CERT_ERROR if r["date"] < K.CERT_DATE.isoformat() else 1.0)
            for r in rows(R.F_OUTLET)}
    diff = np.array([outl[d] - inl[d] - mill[d] for d in sorted(outl) if d in inl and d in mill and d not in
                     ("2025-06-17", "2025-06-18")])
    ok("Z9 two observers of one flow differ like meters: the outlet against the inlets and the mill varies day to day",
       np.std(diff) > 0.02 and 0.0 < np.median(diff) < 1.5, "median %.3f, sd %.3f Ml" % (np.median(diff), np.std(diff)))

    # ---------------------------------------------------------------- goldens
    wb = ox.load_workbook(os.path.join(GOLDEN, DL.F_BOOK), data_only=True)
    s1 = [r for r in wb["Phase 3 2027-28"].iter_rows(min_row=5, max_row=10, values_only=True)]
    ok("Q1 the golden workbook's call block is figures.json's",
       ["%.2f" % r[1] for r in s1[:5]] == [fc["base"], fc["forecast"], fc["reduction"], fc["target"], fc["shortfall"]]
       and s1[5][1] == fc["gang_years"])
    s2 = {r[0]: r[2] for r in wb["Other districts 2025-26"].iter_rows(min_row=5, max_row=16, values_only=True)}
    ok("Q2 the golden's twelve candidates are figures.json's", {c: "%.3f" % v for c, v in s2.items()} == A)
    s3 = [r[1] for r in wb["Zone balance 2025-26"].iter_rows(min_row=5, max_row=10, values_only=True)]
    ok("Q3 the golden's balance is figures.json's", ["%.2f" % v for v in s3] ==
       [B["system_input"], B["billed_metered"], B["billed_unmetered"], B["unbilled"], B["apparent"], B["real_losses"]])
    svg = open(os.path.join(GOLDEN, DL.F_CHART), encoding="utf-8").read()
    ok("Q4 the chart's title states figures.json's reduction, shortfall and target and draws 22 districts",
       "take %s Ml/d" % fc["reduction"] in svg and "%s Ml/d short of the %s" % (fc["shortfall"], fc["target"]) in svg
       and sum(svg.count(c + " ") >= 1 for c in W.PHASE3) == 22)
    ok("Q5 no em dash in the goldens or the write-up",
       "—" not in svg and "—" not in xlsx_text(os.path.join(GOLDEN, DL.F_BOOK))
       and (not os.path.exists(os.path.join(TASK, "submission.md")) or "—" not in open(os.path.join(TASK, "submission.md"), encoding="utf-8").read()))

    n_fail = sum(1 for r, _ in RESULTS if not r)
    print("\n%d checks, %d failed" % (len(RESULTS), n_fail))
    return 1 if n_fail else 0


if __name__ == "__main__":
    sys.exit(main())
