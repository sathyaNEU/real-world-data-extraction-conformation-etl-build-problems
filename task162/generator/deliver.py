"""task162 goldens from the shipped pack: Coldbeck's 2027-28 delivery-plan workbook (the Phase 3 call with the phase's
forecast, and the zone check on its 2025-26 baseline: the twelve other districts' assessed leakage and the zone's
balance), the district chart, and generator/figures.json (every figure the write-up states, as it states it).
Everything is computed by calc.py from the pack, with N1 read off LW-LA-04's Table C2 and the settings off the valve
schedule.

    python3 task162/generator/deliver.py <target> <golden dir> <figures.json>"""
from __future__ import annotations

import calendar
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import calc as C            # noqa: E402
import common as K          # noqa: E402
import postprocess as PP    # noqa: E402
import world as W           # noqa: E402

F_BOOK = "phase3_leakage_forecast_2027_28.xlsx"
F_CHART = "phase3_leakage_by_district.svg"
CREATED, MODIFIED = "2026-11-06T09:12:00Z", "2026-11-09T16:40:00Z"


def r2(x):
    return "%.2f" % K.round_half_up(x, 2)


def r3(x):
    return "%.3f" % K.round_half_up(x, 3)


def export_year(target):
    """Each Phase 3 district's 2025-26 figure as Lynmere's export of 14 May 2026 gives it (Schedule 4.2's baseline)."""
    import openpyxl
    ws = openpyxl.load_workbook(os.path.join(target, "dma_leakage_assessment_phase3_2025_26.xlsx"), read_only=True)["Year"]
    out = {}
    for r in ws.iter_rows(min_row=4, values_only=True):
        if r and r[0] and str(r[0]).startswith("NZ"):
            out[r[0]] = float(r[2])
    return out


def compute(target):
    N1, S = C.read_n1(target), C.read_settings(target)
    P = C.Pack(target, N1, S)
    names = {r["district_code"]: r["district_name"] for r in C._read(os.path.join(target, "district_register.csv"))}
    base_d = export_year(target)
    fc_d = {c: float(P.forecast(c).mean()) for c in W.PHASE3}
    base = sum(base_d[c] for c in W.PHASE3)
    fc = sum(fc_d[c] for c in W.PHASE3)
    red = base - fc
    short = max(0.0, K.TARGET - red)
    gangs = K.ceil_div(short, K.GANG_YIELD) if short > 0 else 0
    rival = C.call(P, W.PHASE3, "nightratio")
    return dict(N1=N1, S=S, names=names, base_d=base_d, fc_d=fc_d, base=base, fc=fc, reduction=red, shortfall=short,
                gangs=gangs, rival=rival, ask_a=C.ask_a(P, W.NEAR), ask_b=C.ask_b(target))


# ------------------------------------------------------------------------------------------------- the workbook
def write_book(path, X):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    navy = "1F3A5F"
    head_fill = PatternFill("solid", fgColor="E3EAF2")
    rule = Border(bottom=Side(style="thin", color="9AA7B5"))
    bold = Font(name="Arial", bold=True, size=10)
    body = Font(name="Arial", size=10)
    small = Font(name="Arial", size=9, color="4A5560")
    wb = openpyxl.Workbook()

    def sheet(ws, widths):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        ws.sheet_view.showGridLines = False

    def title(ws, text, sub):
        ws["A1"] = text
        ws["A1"].font = Font(name="Arial", bold=True, size=13, color=navy)
        ws["A2"] = sub
        ws["A2"].font = small

    def header(ws, row, cols):
        for j, h in enumerate(cols, 1):
            c = ws.cell(row=row, column=j, value=h)
            c.font = bold
            c.fill = head_fill
            c.border = rule
            c.alignment = Alignment(horizontal="left" if j <= 2 else "right", vertical="bottom", wrap_text=True)

    def note(ws, row, text):
        c = ws.cell(row=row, column=1, value=text)
        c.font = small
        return row + 1

    # the call
    ws = wb.active
    ws.title = "Phase 3 2027-28"
    sheet(ws, [58, 12, 12])
    title(ws, "Phase 3 falls %s Ml/d short of its contracted %s Ml/d in 2027-28: add %d find-and-fix gang-years"
          % (r2(X["shortfall"]), r2(K.TARGET), X["gangs"]),
          "Leakage Outcome Contract LW/LOC/2024/07, Schedule 4 (Variation 3): the 22 Phase 3 districts, NZ13 to NZ34, "
          "with fixed-outlet PRVs at the rev C settings from 1 April 2027")
    header(ws, 4, ["", "Value", "Unit"])
    rows = [("2025-26 assessed leakage (Lynmere export of 14 May 2026)", K.round_half_up(X["base"], 2), "Ml/d", "0.00"),
            ("2027-28 forecast leakage", K.round_half_up(X["fc"], 2), "Ml/d", "0.00"),
            ("Reduction", K.round_half_up(X["reduction"], 2), "Ml/d", "0.00"),
            ("Contracted reduction (Schedule 4.3)", K.TARGET, "Ml/d", "0.00"),
            ("Shortfall", K.round_half_up(X["shortfall"], 2), "Ml/d", "0.00"),
            ("Find-and-fix gang-years to add for 2027-28 (Schedule 4.4)", X["gangs"], "gang-years", "0")]
    for i, (lab, v, u, fmt) in enumerate(rows, 5):
        ws.cell(row=i, column=1, value=lab).font = bold if i in (7, 10) else body
        c = ws.cell(row=i, column=2, value=v)
        c.number_format = fmt
        c.font = bold if i in (7, 10) else body
        ws.cell(row=i, column=3, value=u).font = body
    r = 12
    for t in ("Forecast: each district's 2025-26 hourly inlet and average zone pressures replayed through its valve, the "
              "outlet the lower of the inlet pressure and the setting and the zone pressure lowered by the same amount; "
              "net night flow scaled to the replayed AZNP at the district's N1 (LW-LA-04 Table C2), night-day factor "
              "rebuilt from the replayed pressures, day by day, then LW-LA-04's months and year.",
              "On weekdays from about 06:00 to 22:00 the paper mill's draw on the trunk main takes the far districts' "
              "inlets below their settings and the valves stand open. Carrying each valve's night pressure ratio through "
              "the whole day would show a %s Ml/d reduction." % r2(X["rival"]["reduction"]),
              "Conditions as Clive Heath's note of 19 October 2026: each month as recorded in the same month of 2025-26, "
              "no growth. Gang-years at 0.05 Ml/d each, rounded up."):
        c = ws.cell(row=r, column=1, value=t)
        c.font = small
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        ws.row_dimensions[r].height = 40
        r += 1

    # Ask A
    wa = wb.create_sheet("Other districts 2025-26")
    sheet(wa, [10, 20, 14])
    title(wa, "The zone's twelve other districts: 2025-26 assessed leakage by LW-LA-04 issue 6",
          "The districts nearer Thornholme, outside Phase 3, from their own inlet meters and average zone points")
    header(wa, 4, ["District", "Name", "Leakage (Ml/d)"])
    for i, c in enumerate(W.NEAR, 5):
        wa.cell(row=i, column=1, value=c).font = body
        wa.cell(row=i, column=2, value=X["names"][c]).font = body
        v = wa.cell(row=i, column=3, value=K.round_half_up(X["ask_a"][c], 3))
        v.number_format = "0.000"
        v.font = body
    r = 5 + len(W.NEAR) + 1
    r = note(wa, r, "Night use allowed for each dwelling connected that day: a new build counts from its plot connection "
                    "and once, its address record taking over on the same supply point; flats on a shared supply count "
                    "one each.")
    r = note(wa, r, "NZ04: both inlet meters, the second still logged as NZ04A. NZ09 and NZ10 leave out 12 to 20 January "
                    "2026, when NZ09 was fed through boundary valve BV-0910.")
    wa.freeze_panes = "A5"

    # Ask B
    wb_ = wb.create_sheet("Zone balance 2025-26")
    sheet(wb_, [44, 12, 8])
    title(wb_, "Northern zone water balance, 2025-26 (LW-LA-04 section 8)",
          "1 April 2025 to 31 March 2026, megalitres a day over the year's 365 days")
    header(wb_, 4, ["Component", "Ml/d", ""])
    B = X["ask_b"]
    items = [("System input (Thornholme outlet meter TH-OUT-0417)", B["siv"]),
             ("Billed metered consumption", B["metered"]),
             ("Billed unmetered consumption", B["unmeasured"]),
             ("Unbilled authorised use", B["unbilled"]),
             ("Apparent losses", B["apparent"]),
             ("Real losses", B["real"])]
    for i, (lab, v) in enumerate(items, 5):
        wb_.cell(row=i, column=1, value=lab).font = bold if i in (5, 7, 10) else body
        c = wb_.cell(row=i, column=2, value=K.round_half_up(v, 2))
        c.number_format = "0.00"
        c.font = bold if i in (5, 7, 10) else body
    r = 12
    r = note(wb_, r, "System input: daily outlet volumes before 12 November 2025 divided by 0.986 (as found 1.4 per cent "
                     "low, certificate FFS-25-1182).")
    r = note(wb_, r, "Customers: the properties in the zone's 34 districts (NZ04A mapped to NZ04) and the trunk main's "
                     "direct supplies (NZ-TM); bills apportioned to the year by days.")
    r = note(wb_, r, "Flats on a shared supply carry one assessed bill each, so a supply point can show identical lines.")
    r = note(wb_, r, "The 34 districts' assessed leakage sits under these real losses, which also carry the trunk main's own "
                     "losses and the balance's estimates.")

    # notes
    wn = wb.create_sheet("Sources")
    sheet(wn, [34, 92])
    wn["A1"] = "Sources"
    wn["A1"].font = Font(name="Arial", bold=True, size=12, color=navy)
    src = [("Telemetry", "northern_zone_inlet_loggers_2025_26.csv and northern_zone_pressure_loggers_2025_26.csv, "
                         "hourly means in GMT; logger_register.csv for serials, makes and units (1 bar = 10.197 m)"),
           ("Baseline", "dma_leakage_assessment_phase3_2025_26.xlsx, Lynmere export of 14 May 2026"),
           ("Method", "LW-LA-04 issue 6, effective 1 April 2025, with Table C2"),
           ("Valves", "phase3_prv_schedule_rev_C.xlsx, Coldbeck network modelling, 9 October 2026"),
           ("Contract", "outcome_contract_schedule_4_phase3.pdf, Variation 3 of 2 September 2026"),
           ("Properties and bills", "property_extract_2026-11-02.csv and billing_extract_north_2025_26.csv"),
           ("Files taken", "Monday 2 November 2026")]
    for i, (a, b) in enumerate(src, 3):
        wn.cell(row=i, column=1, value=a).font = bold
        wn.cell(row=i, column=2, value=b).font = body
    for w in (ws, wa, wb_):
        w.page_setup.orientation = "portrait"
        w.page_setup.fitToWidth = 1
    wb.save(path)
    PP.normalize_ooxml(path, K.PEOPLE["delivery"], K.PEOPLE["delivery"], CREATED, MODIFIED, excel=True,
                       company=K.PARTNER)


# ------------------------------------------------------------------------------------------------- the chart
def write_chart(path, X):
    Wd, Ht = 880, 820
    left, right, top, bottom = 210, 40, 118, 92
    pw, ph = Wd - left - right, Ht - top - bottom
    xmax = 0.30
    codes = W.PHASE3
    step = ph / len(codes)
    light, dark, line = "#86b6ef", "#1c5cab", "#c3c2b7"
    ink, ink2, muted, grid = "#0b0b0b", "#52514e", "#898781", "#e1e0d9"

    def xp(v):
        return left + pw * v / xmax

    out = ['<svg xmlns="http://www.w3.org/2000/svg" width="%d" height="%d" viewBox="0 0 %d %d" '
           'font-family="Arial, Helvetica, sans-serif">' % (Wd, Ht, Wd, Ht),
           '<rect width="%d" height="%d" fill="#fcfcfb"/>' % (Wd, Ht),
           '<text x="24" y="38" font-size="19" font-weight="bold" fill="%s">Phase 3 valves take %s Ml/d off leakage in '
           '2027-28, %s Ml/d short of the %s contracted</text>' % (ink, r2(X["reduction"]), r2(X["shortfall"]), r2(K.TARGET)),
           '<text x="24" y="62" font-size="13" fill="%s">Assessed leakage by district, 2025-26, and forecast 2027-28 with '
           'the rev C valves in service all year (Ml/d). Phase 3: %s Ml/d to %s Ml/d.</text>'
           % (ink2, r2(X["base"]), r2(X["fc"]))]
    # legend
    out.append('<circle cx="30" cy="90" r="6" fill="%s"/><text x="42" y="94" font-size="12.5" fill="%s">2025-26 assessed '
               '(Lynmere export of 14 May 2026)</text>' % (light, ink))
    out.append('<circle cx="400" cy="90" r="6" fill="%s"/><text x="412" y="94" font-size="12.5" fill="%s">2027-28 '
               'forecast</text>' % (dark, ink))
    for k in range(0, 7):
        v = 0.05 * k
        x = xp(v)
        out.append('<line x1="%.1f" y1="%d" x2="%.1f" y2="%d" stroke="%s" stroke-width="1"/>' % (x, top - 6, x, top + ph, grid))
        out.append('<text x="%.1f" y="%d" font-size="11.5" fill="%s" text-anchor="middle">%.2f</text>'
                   % (x, top + ph + 18, muted, v))
    out.append('<text x="%.1f" y="%d" font-size="12" fill="%s" text-anchor="middle">Leakage, megalitres a day</text>'
               % (left + pw / 2, top + ph + 38, ink2))
    for i, c in enumerate(codes):
        y = top + step * (i + 0.5)
        b, f = X["base_d"][c], X["fc_d"][c]
        out.append('<text x="%d" y="%.1f" font-size="12" fill="%s" text-anchor="end">%s %s</text>'
                   % (left - 14, y + 4, ink, c, X["names"][c]))
        out.append('<line x1="%.1f" y1="%.1f" x2="%.1f" y2="%.1f" stroke="%s" stroke-width="2"/>'
                   % (xp(f), y, xp(b), y, line))
        out.append('<circle cx="%.1f" cy="%.1f" r="5.5" fill="%s" stroke="#fcfcfb" stroke-width="2"/>' % (xp(b), y, light))
        out.append('<circle cx="%.1f" cy="%.1f" r="5.5" fill="%s" stroke="#fcfcfb" stroke-width="2"/>' % (xp(f), y, dark))
    out.append('<text x="24" y="%d" font-size="11" fill="%s">Districts in order along the Thornholme to Brackenby trunk '
               'main. Forecast: 2025-26 hourly pressures replayed through each valve\'s setting, LW-LA-04 issue 6.</text>'
               % (Ht - 30, ink2))
    out.append('<text x="24" y="%d" font-size="11" fill="%s">Coldbeck Network Partners, Northern zone delivery, '
               '9 November 2026</text>' % (Ht - 13, muted))
    out.append("</svg>")
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write("\n".join(out) + "\n")


# ------------------------------------------------------------------------------------------------- figures.json
def figures(X):
    B = X["ask_b"]
    return {
        "call": {"meets_target": "no", "base": r2(X["base"]), "forecast": r2(X["fc"]), "reduction": r2(X["reduction"]),
                 "target": r2(K.TARGET), "shortfall": r2(X["shortfall"]), "gang_years": X["gangs"],
                 "gang_yield": r2(K.GANG_YIELD), "default_reduction": r2(X["rival"]["reduction"]),
                 "districts": len(W.PHASE3)},
        "steps": {"bar_m": "10.197", "cert_divisor": "%.3f" % (1.0 + K.CERT_ERROR), "cert_pct": "1.4",
                  "near_districts": len(W.NEAR), "zone_districts": len(W.PHASE3) + len(W.NEAR)},
        "answers": {"ask_a": {c: r3(X["ask_a"][c]) for c in W.NEAR},
                    "ask_b": {"system_input": r2(B["siv"]), "billed_metered": r2(B["metered"]),
                              "billed_unmetered": r2(B["unmeasured"]), "unbilled": r2(B["unbilled"]),
                              "apparent": r2(B["apparent"]), "real_losses": r2(B["real"])}},
        "unrounded": {"base": round(X["base"], 5), "forecast": round(X["fc"], 5), "reduction": round(X["reduction"], 5),
                      "shortfall": round(X["shortfall"], 5), "default_reduction": round(X["rival"]["reduction"], 5),
                      "ask_a": {c: round(X["ask_a"][c], 5) for c in W.NEAR},
                      "ask_b": {k: round(v, 5) for k, v in B.items()}},
    }


def main(target, golden, fig_path):
    os.makedirs(golden, exist_ok=True)
    for f in os.listdir(golden):
        if f not in (F_BOOK, F_CHART):
            os.remove(os.path.join(golden, f))
    X = compute(target)
    write_book(os.path.join(golden, F_BOOK), X)
    write_chart(os.path.join(golden, F_CHART), X)
    with open(fig_path, "w", encoding="utf-8") as fh:
        json.dump(figures(X), fh, indent=2, ensure_ascii=False, sort_keys=True)
        fh.write("\n")
    mt = calendar.timegm((2026, 11, 9, 16, 40, 0, 0, 0, 0))
    for f in os.listdir(golden):
        os.utime(os.path.join(golden, f), (mt, mt))
    return X


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3])
