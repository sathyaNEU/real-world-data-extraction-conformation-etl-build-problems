"""task162 workbooks in the pack: Lynmere's 2025-26 assessment export for the Phase 3 districts and Coldbeck's Phase 3
valve schedule (rev C)."""
from __future__ import annotations

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import common as K
import docs as DOC

HEAD_FILL = PatternFill("solid", fgColor="DCE6EE")
THIN = Side(style="thin", color="B8C1C9")


def _head(ws, row, cols, widths):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = Font(bold=True, name="Calibri", size=10)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = Border(bottom=THIN)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


EXPORT_COLS = ["District", "Month", "Valid days", "Households (mean)", "Non-households (mean)", "MNF (l/s)",
               "Night use allowance (l/s)", "Net night flow (l/s)", "AZNP (m)", "Night-day factor (h)",
               "Leakage (Ml/d)"]


def write_export(path, monthly, yearly):
    """monthly: rows [code, 'YYYY-MM', valid, hh, nhh, mnf, allow, nnf, aznp, ndf, leak];
    yearly: rows [code, name, leak_year]."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Monthly"
    ws.cell(row=1, column=1, value="Lynmere Water: DMA leakage assessment, 2025-26, Phase 3 districts (LW-LA-04 issue 6)")
    ws.cell(row=1, column=1).font = Font(bold=True, size=11)
    ws.cell(row=2, column=1, value="Export of 14 May 2026, %s, Leakage Analysis" % K.PEOPLE["leakage"])
    _head(ws, 4, EXPORT_COLS, [9, 9, 7, 11, 11, 9, 11, 10, 9, 10, 10])
    for i, r in enumerate(monthly, 5):
        vals = [r[0], r[1], r[2], round(r[3], 1), round(r[4], 1), round(r[5], 2), round(r[6], 3), round(r[7], 3),
                round(r[8], 1), round(r[9], 2), round(r[10], 4)]
        for j, v in enumerate(vals, 1):
            ws.cell(row=i, column=j, value=v)
    ws.freeze_panes = "C5"
    wy = wb.create_sheet("Year")
    wy.cell(row=1, column=1, value="2025-26 assessed leakage, mean of the twelve months (Ml/d)")
    wy.cell(row=1, column=1).font = Font(bold=True, size=11)
    _head(wy, 3, ["District", "Name", "Leakage 2025-26 (Ml/d)"], [9, 20, 12])
    for i, r in enumerate(yearly, 4):
        wy.cell(row=i, column=1, value=r[0])
        wy.cell(row=i, column=2, value=r[1])
        wy.cell(row=i, column=3, value=round(r[2], 4))
    last = 3 + len(yearly)
    wy.cell(row=last + 1, column=2, value="Phase 3")
    wy.cell(row=last + 1, column=3, value="=SUM(C4:C%d)" % last)
    wy.cell(row=last + 1, column=2).font = Font(bold=True)
    wb.save(path)


def write_valves(path, rows):
    """rows: [code, name, size, setting, aznp_before, aznp_after, week]"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Phase 3 PRVs"
    ws.cell(row=1, column=1, value="Coldbeck Network Partners: Northern zone Phase 3 pressure reducing valves, schedule rev C")
    ws.cell(row=1, column=1).font = Font(bold=True, size=11)
    ws.cell(row=2, column=1, value="Prepared by %s, network modelling, 9 October 2026" % K.PEOPLE["modeller"])
    ws.cell(row=3, column=1, value=DOC.VALVE_BASIS)
    ws.cell(row=3, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=8)
    ws.row_dimensions[3].height = 54
    _head(ws, 5, DOC.VALVE_COLS, [9, 20, 8, 12, 12, 13, 15, 14])
    for i, r in enumerate(rows, 6):
        vals = [r[0], r[1], "DN%d" % r[2], "fixed outlet", r[3], round(r[4], 1), round(r[5], 1), r[6]]
        for j, v in enumerate(vals, 1):
            ws.cell(row=i, column=j, value=v)
    ws.freeze_panes = "C6"
    wb.save(path)
