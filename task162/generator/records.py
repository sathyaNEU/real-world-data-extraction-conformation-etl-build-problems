"""task162 data files cut from the truth: the old tracker's export, the public data portal's daily version shares and
daily active installs, the old triage dashboard, the public bug statistics, the stability series, the new tracker's
parallel-run export, the component-group register and the triage rota."""
from __future__ import annotations

import datetime as dt
import math
import re

import numpy as np

import common as K
import parallel as PR
import world as W

F_EXPORT = "old_tracker_export_20251001_20261217.csv"
F_FIELDS = "old_tracker_export_fields.txt"
F_SHARES = "version_share_daily.csv"
F_INSTALLS = "active_installs_daily.csv"
F_DASH = "triage_dashboard_weekly.xlsx"
F_PUBLIC = "public_bug_statistics_monthly.csv"
F_STABILITY = "stability_monthly.csv"
F_NEWEXPORT = "newtracker_parallel_run_export.json"
F_GROUPS = "component_groups.xlsx"
F_ROTA = "triage_rota_2026-11.xlsx"

EXPORT_COLS = ["bug_id", "opened", "changeddate", "product", "component", "version", "op_sys", "bug_severity",
               "bug_status", "resolution", "dupe_of", "groups", "keywords"]
SHARE_COLS = ["date", "platform", "version", "share_pct"]
INSTALL_COLS = ["date", "platform", "active_installs_thousands"]
PUBLIC_COLS = ["month", "platform", "reports_filed", "reports_per_100k_active_installs"]
STABILITY_COLS = ["month", "platform", "version", "usage_hours_millions", "crash_reports", "crashes_per_1000_hours"]
PULL = K.PULL
BULK_CLOSE = dt.datetime(2026, 11, 30, 9, 14, 0)
ARCHIVE = ("Archive", "Parallel run 2026-11")


def fmt(t):
    return t.strftime("%Y-%m-%d %H:%M:%S")


def old_states(T, reps):
    """Each exported report's state on the old tracker at the pull: component, status, resolution, dupe_of and last
    change. The parallel run's reports were worked on the new tracker and closed MOVED in bulk on 30 November."""
    r = K.rng("old_states")
    out = []
    run_a = PR.RUN_FROM
    run_b = PR.RUN_TO
    for x in reps:
        if x["kind"] not in ("baseline", "first_use"):
            continue
        d = x["opened"].date()
        if d < K.RECORD_START or d > K.AS_OF:
            continue
        s = dict(x)
        filed = x["opened"]
        default = PR.form_default(x["product"], x["component"])
        triage_at = filed + dt.timedelta(hours=float(r.lognormal(math.log(14), 0.9)))
        if run_a <= filed < run_b:
            # closed MOVED in bulk after the run and moved to the archive product
            s["component_now"] = ARCHIVE
            s["status"], s["resolution"], s["dupe"] = "RESOLVED", "MOVED", None
            s["changed"] = BULK_CLOSE + dt.timedelta(seconds=int((x["bug_id"] % 97) * 3 + r.integers(0, 3)))
            s["keywords"] = ""
            out.append(s)
            continue
        comp = (x["product"], x["component"])
        if x["moved_webcompat"]:
            comp = W.WEBCOMPAT_PRODUCT
        last = filed
        status, res, dupe = "UNCONFIRMED", "", None
        if triage_at < PULL:
            status = "NEW"
            last = triage_at
            s["component_now"] = comp
        else:
            s["component_now"] = default
        if x["dupe_of"] is not None:
            at = filed + dt.timedelta(hours=float(r.lognormal(math.log(40), 0.8)))
            if at < PULL:
                status, res, dupe, last = "RESOLVED", "DUPLICATE", x["dupe_of"], at
                s["component_now"] = comp
        else:
            at = filed + dt.timedelta(days=float(r.lognormal(math.log(11), 1.3)))
            if at < PULL and r.random() < 0.74:
                res = str(r.choice(["FIXED", "INCOMPLETE", "WORKSFORME", "INVALID", "WONTFIX"],
                                   p=[0.30, 0.26, 0.17, 0.15, 0.12]))
                status, last = ("VERIFIED" if res == "FIXED" and r.random() < 0.18 else "RESOLVED"), at
                s["component_now"] = comp
            elif status == "NEW" and r.random() < 0.22:
                status = "ASSIGNED"
                last = triage_at + dt.timedelta(hours=float(r.uniform(2, 200)))
                if last >= PULL:
                    last = triage_at
        s["status"], s["resolution"], s["dupe"] = status, res, dupe
        s["changed"] = max(last, filed + dt.timedelta(seconds=40))
        if x["regression_key"] is not None and x["dupe_of"] is None:
            s["keywords"] = "regression" if r.random() < 0.62 else ""
        elif x["product"] == "Install and Update" and x["component"] == "Crash Reporter":
            s["keywords"] = "crash"
        else:
            s["keywords"] = ""
        out.append(s)
    return out


def export_rows(states):
    for s in sorted(states, key=lambda s: s["bug_id"]):
        yield {"bug_id": s["bug_id"], "opened": fmt(s["opened"]), "changeddate": fmt(s["changed"]),
               "product": s["component_now"][0], "component": s["component_now"][1], "version": s["version"],
               "op_sys": s["op_sys"], "bug_severity": s["severity"], "bug_status": s["status"],
               "resolution": s["resolution"] or "---", "dupe_of": s["dupe"] if s["dupe"] else "",
               "groups": "core-security" if s["security"] else "", "keywords": s["keywords"]}


# ------------------------------------------------------------------------------------------------- the portal
def share_rows(T):
    """Daily share of each platform's active installs on each release-channel major version, as the portal publishes
    it (per cent to two decimals; majors under 0.05 per cent folded into 'older')."""
    r = K.rng("portal_shares")
    for d in W.DAYS:
        if d < K.RECORD_START or d > K.AS_OF:
            continue
        i = W.IDX[d]
        for p in W.PLATFORMS:
            st = {}
            for v, sh in T.adopt.share[p][i].items():
                m = int(v.split(".")[0])
                st[m] = st.get(m, 0.0) + sh
            vals = {}
            for m, sh in st.items():
                vals[m] = max(0.0, 100 * sh + float(r.normal(0, 0.035)) * (1 if sh > 0.002 else 0))
            keep = {m: s for m, s in vals.items() if s >= 0.05}
            older = sum(s for m, s in vals.items() if s < 0.05)
            tot = sum(keep.values()) + older
            for m in sorted(keep):
                yield {"date": d.isoformat(), "platform": p, "version": "%d" % m,
                       "share_pct": "%.2f" % (keep[m] * 100 / tot)}
            if older >= 0.005:
                yield {"date": d.isoformat(), "platform": p, "version": "older", "share_pct": "%.2f" % (older * 100 / tot)}


def install_rows(T):
    r = K.rng("portal_installs")
    for d in W.DAYS:
        if d < K.RECORD_START or d > K.AS_OF:
            continue
        for p in W.PLATFORMS:
            v = W.dai(d, p) * (1 + float(r.normal(0, 0.0022)))
            yield {"date": d.isoformat(), "platform": p, "active_installs_thousands": int(round(v / 1000))}


# ------------------------------------------------------------------------------------------------- the dashboard
DASH_PRODUCTS = ["Engine", "Shell", "Extensions", "DevTools", "Install and Update", "Web Compatibility"]
DASH_EXCLUDED = ("DUPLICATE", "INVALID", "WORKSFORME", "MOVED")
DASH_FIRST = dt.date(2025, 9, 29)
DASH_LAST = dt.date(2026, 12, 7)


def dashboard_rows(states):
    """New defects by week filed and product, the old dashboard's basis, at the pull."""
    weeks = []
    m = DASH_FIRST
    while m <= DASH_LAST:
        weeks.append(m)
        m += dt.timedelta(days=7)
    cnt = {w: {p: 0 for p in DASH_PRODUCTS} for w in weeks}
    for s in states:
        if s["security"] or s["resolution"] in DASH_EXCLUDED:
            continue
        w = K.monday(s["opened"].date())
        if w in cnt:
            cnt[w][s["component_now"][0]] += 1
    return [(w, cnt[w]) for w in weeks]


def write_dashboard(path, states):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Weekly"
    ws["A1"] = "Triage dashboard: new defects by week filed"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Reports filed through the report form in the week (Monday to Sunday, Helsinki time), by product they "
                "are in now. The security group is left out, and so are reports resolved DUPLICATE, INVALID, "
                "WORKSFORME or MOVED.")
    ws["A3"] = "Old tracker, refreshed 18.12.2026 06:40 with the weekly export. Dashboard v3.2, triage team."
    ws["A2"].alignment = Alignment(wrap_text=False)
    head = ["Week starting"] + DASH_PRODUCTS + ["New defects"]
    for j, h in enumerate(head):
        c = ws.cell(row=5, column=j + 1, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCE6F0")
    rows = dashboard_rows(states)
    for i, (w, c) in enumerate(rows):
        rr = 6 + i
        ws.cell(row=rr, column=1, value=w).number_format = "d.m.yyyy"
        for j, p in enumerate(DASH_PRODUCTS):
            ws.cell(row=rr, column=2 + j, value=c[p])
        ws.cell(row=rr, column=2 + len(DASH_PRODUCTS), value="=SUM(B%d:G%d)" % (rr, rr))
    last = 6 + len(rows)
    ws.cell(row=last + 1, column=1, value="29.9.2025: the record starts on Wednesday 1 October 2025.")
    ws.cell(row=last + 2, column=1, value=("2.11. to 23.11.2026: the parallel run's reports were triaged on the new "
                                           "tracker and closed MOVED here on 30.11.2026."))
    ws.column_dimensions["A"].width = 14
    for j in range(len(DASH_PRODUCTS) + 1):
        ws.column_dimensions["BCDEFGH"[j]].width = 13
    ws.freeze_panes = "B6"
    wb.save(path)
    return rows


# ------------------------------------------------------------------------------------------------- public statistics
def months(a, b):
    out = []
    y, m = a.year, a.month
    while (y, m) <= (b.year, b.month):
        out.append((y, m))
        m += 1
        if m == 13:
            y, m = y + 1, 1
    return out


def public_rows(states, T):
    """The vendor's published monthly bug statistics: reports filed through the form by platform (security reports
    not published), and per 100,000 of the month's average daily active installs."""
    cnt = {}
    for s in states:
        if s["security"]:
            continue
        p = "Windows" if s["op_sys"].startswith("Windows") else "macOS" if s["op_sys"].startswith("macOS") else "Linux"
        key = (s["opened"].year, s["opened"].month, p)
        cnt[key] = cnt.get(key, 0) + 1
    for (y, m) in months(K.RECORD_START, dt.date(2026, 11, 1)):
        days = [d for d in W.DAYS if d.year == y and d.month == m and d >= K.RECORD_START]
        for p in W.PLATFORMS:
            avg = sum(W.dai(d, p) for d in days) / len(days)
            n = cnt.get((y, m, p), 0)
            yield {"month": "%d-%02d" % (y, m), "platform": p, "reports_filed": n,
                   "reports_per_100k_active_installs": "%.1f" % K.round_half_up(n / avg * 1e5, 1)}


# ------------------------------------------------------------------------------------------------- stability
def stability_rows(T):
    """The crash reporter's monthly stability by release-channel major version and platform."""
    r = K.rng("stability")
    rate = {}
    for m, _ in K.RELEASES:
        for p in W.PLATFORMS:
            rate[(m, p)] = {"Windows": 0.92, "macOS": 0.61, "Linux": 0.74}[p] * float(r.lognormal(0, 0.18))
    hours_per_day = {"Windows": 2.71, "macOS": 2.38, "Linux": 2.95}
    for (y, mo) in months(K.RECORD_START, dt.date(2026, 11, 1)):
        days = [d for d in W.DAYS if d.year == y and d.month == mo and d >= K.RECORD_START]
        for p in W.PLATFORMS:
            use = {}
            for d in days:
                st = T.adopt.share[p][W.IDX[d]]
                for v, sh in st.items():
                    major = int(v.split(".")[0])
                    use[major] = use.get(major, 0.0) + W.dai(d, p) * sh * hours_per_day[p]
            tot = sum(use.values())
            for major in sorted(use):
                if use[major] < 0.01 * tot:
                    continue
                crashes = int(r.poisson(use[major] / 1000 * rate.get((major, p), rate[(31, p)])))
                yield {"month": "%d-%02d" % (y, mo), "platform": p, "version": "%d" % major,
                       "usage_hours_millions": "%.2f" % (use[major] / 1e6), "crash_reports": crashes,
                       "crashes_per_1000_hours": "%.2f" % (crashes / use[major] * 1000)}


# ------------------------------------------------------------------------------------------------- the new tracker
STATUS_ID = {"New": "10000", "Triaged": "10101", "Resolved": "5", "Closed": "6", "Needs diagnosis": "10411",
             "Diagnosed": "10412", "Done": "10001"}
STATUS_CAT = {"New": ("new", "To Do"), "Triaged": ("indeterminate", "In Progress"), "Resolved": ("done", "Done"),
              "Closed": ("done", "Done"), "Needs diagnosis": ("new", "To Do"), "Diagnosed": ("indeterminate", "In Progress"),
              "Done": ("done", "Done")}
RESOLUTION_ID = {"Fixed": "1", "Won't Fix": "2", "Duplicate": "3", "Incomplete": "4", "Cannot Reproduce": "5",
                 "Invalid": "10100", "Done": "10000"}
PRIORITY_ID = {"P1": "1", "P2": "2", "P3": "3", "P4": "4", "Unprioritised": "10600"}
PROJECT = {"BRW": ("10201", "Browser intake"), "WEBC": ("10305", "Web compatibility")}
WORKFLOW = {"BRW": "BRW: Triage workflow", "WEBC": "WEBC: Site report workflow"}
LINK_TYPE = {"id": "10002", "name": "Duplicate", "inward": "is duplicated by", "outward": "duplicates"}


def version_key(v):
    m = re.match(r"(\d+)\.(\d+)(?:\.(\d+)|b(\d+)|a(\d+))?$", v)
    major, minor = int(m.group(1)), int(m.group(2))
    if m.group(4):
        return (major, minor, 1, int(m.group(4)))
    if m.group(5):
        return (major, minor, 0, int(m.group(5)))
    return (major, minor, 2, int(m.group(3) or 0))


def jtime(t_local, ms):
    u = K.to_utc(t_local)
    return u.strftime("%Y-%m-%dT%H:%M:%S") + ".%03d+0000" % ms


def newtracker(issues):
    """The migration team's export of the run: every issue created in BRW and WEBC from 2 November to the bridge's
    switch-off on 30 November, with its changelog, as the REST search returns it (fields listed, times UTC)."""
    r = K.rng("jira_export")
    comps = sorted({it["comp0"] for it in issues} | set(PR.NEW_COMPONENT.values()) | {PR.WEBC_COMPONENT})
    comp_id = {c: str(10230 + k) for k, c in enumerate(comps)}
    vers = sorted({it["src"]["version"] for it in issues if it["src"]}, key=version_key)
    ver_id = {v: str(11690 + k) for k, v in enumerate(vers)}
    order = sorted(issues, key=lambda it: (it["created"], it["legacy"] or 0, it["copy"]))
    iid = {id(it): str(1184033 + k) for k, it in enumerate(order)}
    # link entries on the issue linked to (Jira writes the link on both ends)
    inward = {}
    for it in issues:
        for h in it["hist"]:
            for f, a, b in h["items"]:
                if f == "Link" and b[1][0] == "issue":
                    inward.setdefault(id(b[1][1]), []).append((h, it))
    hist_all = []
    for it in issues:
        for h in it["hist"]:
            hist_all.append((h["at"], 0, id(it), h, it))
        for h, src in inward.get(id(it), []):
            hist_all.append((h["at"], 1, id(it), h, src))
    hist_all.sort(key=lambda e: (e[0], e[1], iid[e[2]]))
    hid = {}
    for k, e in enumerate(hist_all):
        hid[(e[1], e[2], id(e[3]))] = str(2840117 + k)
    ms_of = {}

    def ms(key):
        if key not in ms_of:
            ms_of[key] = int(r.integers(0, 1000))
        return ms_of[key]

    link_ids = {}

    def link_id(a, b):
        k = (a, b)
        if k not in link_ids:
            link_ids[k] = str(70311 + 3 * len(link_ids) + int(r.integers(0, 3)))
        return link_ids[k]

    out = []
    for it in sorted(issues, key=lambda it: (it["project"] if not it.get("moved_at") else "WEBC",
                                             int(PR.current_key(it).split("-")[1]))):
        x = it["src"]
        status, resolution, resdate, prio = "New", None, None, "Unprioritised"
        if it["project"] == "WEBC":
            status = "Needs diagnosis"
        comps_now = [it["comp0"]]
        legacy_now = None
        old_id_now = None
        updated = it["created"]
        links = []
        histories = []
        entries = [(h, 0, it) for h in it["hist"]] + [(h, 1, src) for h, src in inward.get(id(it), [])]
        entries.sort(key=lambda e: (e[0]["at"], e[1]))
        for h, kind, src in entries:
            items = []
            at = h["at"]
            if kind == 1:
                tgt_key = PR.key_at(src, at)
                items.append({"field": "Link", "fieldtype": "jira", "from": None, "fromString": None, "to": tgt_key,
                              "toString": "This issue is duplicated by %s" % tgt_key})
                links.append({"id": link_id(id(src), id(it)), "type": dict(LINK_TYPE),
                              "inwardIssue": {"id": iid[id(src)], "key": PR.current_key(src)}})
            else:
                for f, a, b in h["items"]:
                    if f == "Legacy bug":
                        items.append({"field": "Legacy bug", "fieldtype": "custom", "fieldId": "customfield_10410",
                                      "from": None, "fromString": a, "to": None, "toString": b})
                        legacy_now = b
                    elif f == "Old tracker ID":
                        items.append({"field": "Old tracker ID", "fieldtype": "custom", "fieldId": "customfield_10522",
                                      "from": None, "fromString": a, "to": None, "toString": b})
                        old_id_now = b
                    elif f == "Affects Version/s":
                        items.append({"field": "Version", "fieldtype": "jira", "fieldId": "versions", "from": None,
                                      "fromString": None, "to": ver_id[b], "toString": b})
                    elif f == "status":
                        items.append({"field": "status", "fieldtype": "jira", "fieldId": "status",
                                      "from": STATUS_ID[status], "fromString": status, "to": STATUS_ID[b],
                                      "toString": b})
                        status = b
                    elif f == "resolution":
                        items.append({"field": "resolution", "fieldtype": "jira", "fieldId": "resolution",
                                      "from": None, "fromString": None, "to": RESOLUTION_ID[b], "toString": b})
                        resolution, resdate = b, at
                    elif f == "priority":
                        items.append({"field": "priority", "fieldtype": "jira", "fieldId": "priority",
                                      "from": PRIORITY_ID[prio], "fromString": prio, "to": PRIORITY_ID[b], "toString": b})
                        prio = b
                    elif f == "Component":
                        if a is not None:
                            items.append({"field": "Component", "fieldtype": "jira", "fieldId": "components",
                                          "from": comp_id[a], "fromString": a, "to": None, "toString": None})
                            comps_now.remove(a)
                        else:
                            items.append({"field": "Component", "fieldtype": "jira", "fieldId": "components",
                                          "from": None, "fromString": None, "to": comp_id[b], "toString": b})
                            comps_now.append(b)
                    elif f == "project":
                        items.append({"field": "project", "fieldtype": "jira", "fieldId": "project",
                                      "from": PROJECT["BRW"][0], "fromString": PROJECT["BRW"][1],
                                      "to": PROJECT["WEBC"][0], "toString": PROJECT["WEBC"][1]})
                        items.append({"field": "Workflow", "fieldtype": "jira", "from": None,
                                      "fromString": WORKFLOW["BRW"], "to": None, "toString": WORKFLOW["WEBC"]})
                    elif f == "Key":
                        items.append({"field": "Key", "fieldtype": "jira", "from": None, "fromString": it["key0"],
                                      "to": None, "toString": it["key_moved"]})
                    elif f == "Link":
                        tgt = b[1]
                        if tgt[0] == "issue":
                            k2 = PR.key_at(tgt[1], at)
                            links.append({"id": link_id(id(it), id(tgt[1])), "type": dict(LINK_TYPE),
                                          "outwardIssue": {"id": iid[id(tgt[1])], "key": PR.current_key(tgt[1])}})
                        else:
                            k2 = it["preload_key"]
                            links.append({"id": link_id(id(it), k2), "type": dict(LINK_TYPE),
                                          "outwardIssue": {"id": str(1100000 + int(k2.split("-")[1]) * 3 % 83000),
                                                           "key": k2}})
                        items.append({"field": "Link", "fieldtype": "jira", "from": None, "fromString": None,
                                      "to": k2, "toString": "This issue duplicates %s" % k2})
            author = h["author"]
            histories.append({"id": hid[(kind, id(it), id(h))],
                              "author": {"displayName": author, "active": True,
                                         "timeZone": "Europe/Helsinki" if author != PR.BRIDGE else "Etc/UTC"},
                              "created": jtime(at, ms(("h", kind, id(it), id(h)))), "items": items})
            updated = max(updated, at)
        if it["issuetype"] == "Task":
            reporter = it["task_by"]
        else:
            reporter = PR.BRIDGE
        proj = "WEBC" if (it["project"] == "WEBC" or it.get("moved_at")) else "BRW"
        sc = STATUS_CAT[status]
        fields = {
            "issuetype": {"name": it["issuetype"], "subtask": False},
            "project": {"id": PROJECT[proj][0], "key": proj, "name": PROJECT[proj][1]},
            "status": {"id": STATUS_ID[status], "name": status, "statusCategory": {"key": sc[0], "name": sc[1]}},
            "resolution": None if resolution is None else {"id": RESOLUTION_ID[resolution], "name": resolution},
            "resolutiondate": None if resdate is None else jtime(resdate, ms(("r", id(it)))),
            "created": jtime(it["created"], ms(("c", id(it)))),
            "updated": jtime(updated, ms(("u", id(it)))) if updated != it["created"] else jtime(it["created"], ms(("c", id(it)))),
            "priority": {"id": PRIORITY_ID[prio], "name": prio},
            "components": [{"id": comp_id[c], "name": c} for c in comps_now],
            "versions": [{"id": ver_id[x["version"]], "name": x["version"]}] if x and not it["bulk"] else [],
            "labels": ["legacy-%d" % it["legacy"]] if it["bulk"] and it["legacy"] is not None else [],
            "environment": x["op_sys"] if x else None,
            "security": {"id": "10050", "name": "Security team"} if it["restricted"] else None,
            "reporter": {"displayName": reporter, "active": True},
            "issuelinks": links,
            "customfield_10410": float(legacy_now) if legacy_now is not None else None,
            "customfield_10522": old_id_now,
        }
        out.append({"expand": "operations,versionedRepresentations,editmeta,changelog,renderedFields",
                    "id": iid[id(it)], "key": PR.current_key(it), "fields": fields,
                    "changelog": {"startAt": 0, "maxResults": len(histories), "total": len(histories),
                                  "histories": histories}})
    names = {"issuetype": "Issue Type", "project": "Project", "status": "Status", "resolution": "Resolution",
             "resolutiondate": "Resolved", "created": "Created", "updated": "Updated", "priority": "Priority",
             "components": "Component/s", "versions": "Affects Version/s", "environment": "Environment",
             "security": "Security Level", "reporter": "Reporter", "issuelinks": "Linked Issues", "labels": "Labels",
             "customfield_10410": "Legacy bug", "customfield_10522": "Old tracker ID"}
    return {"expand": "names,schema", "startAt": 0, "maxResults": len(out), "total": len(out), "issues": out,
            "names": names}


def task_history(issues):
    """The team's own tasks: housekeeping on the new tracker, most done within the run."""
    r = K.rng("tasks")
    for it in issues:
        if it["issuetype"] != "Task":
            continue
        if r.random() < 0.7:
            at = PR.next_working(it["created"] + dt.timedelta(hours=float(r.uniform(2, 120))))
            if at is not None:
                it["hist"].append(dict(author=it["task_by"], at=at,
                                       items=[("status", "New", "Done"), ("resolution", None, "Done")]))


# ------------------------------------------------------------------------------------------------- register and rota
REGISTER_LEADS = {"Rendering and Layout": "Johannes Laukkanen", "Graphics and Media": "Saara Väisänen",
                  "Networking": "Tapio Kallio", "Script and Web APIs": "Johannes Laukkanen",
                  "Browser Interface": "Linnea Kärki", "Extensions": "Arja Turunen",
                  "Developer Tools": "Sakari Puhakka", "Platform and Updates": "Tapio Kallio"}


def register_rows():
    rows = []
    for g in PR.GROUPS:
        for c in sorted(set(PR.NEW_COMPONENT.values())):
            if PR.group_of(c) == g:
                rows.append((g, c, REGISTER_LEADS[g]))
    return rows


def write_groups(path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Component groups"
    ws["A1"] = "Browser intake (BRW): component groups"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Each group is one area rota on the new tracker. A report with components in more than one group is "
                "in each of those groups' queues.")
    for j, h in enumerate(["Group", "Component", "Area contact"]):
        c = ws.cell(row=4, column=j + 1, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E2EFDA")
    rows = register_rows()
    for i, (g, c, lead) in enumerate(rows):
        ws.cell(row=5 + i, column=1, value=g)
        ws.cell(row=5 + i, column=2, value=c)
        ws.cell(row=5 + i, column=3, value=lead)
    n = 5 + len(rows)
    ws.cell(row=n + 1, column=1, value="Not in a group")
    ws.cell(row=n + 1, column=2, value=PR.HOUSEKEEPING)
    ws.cell(row=n + 1, column=3, value="the triage team's own tasks")
    ws.cell(row=n + 3, column=1, value="Agreed in the triage team, 26.10.2026 (UK).")
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 26
    wb.save(path)
    return rows


def rota_rows():
    rows = []
    d = K.RUN_START
    while d <= dt.date(2026, 12, 4):
        wk = K.monday(d)
        names, sat = PR.ROTA[wk]
        if d.weekday() < 5:
            rows.append((d, "08:00-16:30", ", ".join(names), "follow-up" if d > K.RUN_END else ""))
        elif d.weekday() == 5 and sat:
            rows.append((d, "10:00-14:00", sat, "Saturday cover"))
        d += dt.timedelta(days=1)
    return rows


def write_rota(path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rota"
    ws["A1"] = "Triage rota, parallel run 2.11. to 29.11.2026 and the follow-up week"
    ws["A1"].font = Font(bold=True, size=13)
    for j, h in enumerate(["Date", "Day", "Hours", "On duty", "Note"]):
        c = ws.cell(row=3, column=j + 1, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FCE4D6")
    for i, (d, hours, who, note) in enumerate(rota_rows()):
        ws.cell(row=4 + i, column=1, value=d).number_format = "d.m.yyyy"
        ws.cell(row=4 + i, column=2, value=["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"][d.weekday()])
        ws.cell(row=4 + i, column=3, value=hours)
        ws.cell(row=4 + i, column=4, value=who)
        ws.cell(row=4 + i, column=5, value=note or None)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 6
    ws.column_dimensions["C"].width = 13
    ws.column_dimensions["D"].width = 72
    ws.column_dimensions["E"].width = 16
    wb.save(path)
