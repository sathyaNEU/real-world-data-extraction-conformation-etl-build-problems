"""The golden, computed from the pack alone: each region's intergenerational elasticity on the parents' long-run
income (every child, the census year's spread less the transitory variance measured within the five-year history),
the call, and the Bridges asks (young people started in each region, following transfers and the April 2021 import
notes; spending in each region, the December accruals taken back from the offices and the central costs shared by
starts). Then the three deliverables."""
from __future__ import annotations

import csv
import datetime as dt
import os

import numpy as np

import common as K

AS_NAMES = dict(zip(K.CODES, K.NAMES))


def read(folder, name):
    with open(os.path.join(folder, name), encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def load(folder):
    cpi = {int(r["year"]): float(r["cpi_2005_100"]) for r in read(folder, K.CPI_FILE)}
    recode = {r["census_region_2001"]: r["region_2021"] for r in read(folder, K.RECODE)}
    cohort = read(folder, K.COHORT)
    earn = {}
    for r in read(folder, K.EARN):
        earn.setdefault(r["child_id"], {})[int(r["tax_year"])] = float(r["earnings"])
    hist = {}
    for r in read(folder, K.HIST):
        hist.setdefault(r["family_id"], {})[int(r["tax_year"])] = float(r["family_total_income_2005_dollars"])
    return dict(cpi=cpi, recode=recode, cohort=cohort, earn=earn, hist=hist)


def children(D):
    """The brief's children: parents' census income above zero, earnings in at least one year at 30 to 34; the log of
    the average yearly earnings at 30 to 34 in 2005 dollars (zeros counted) and the log of the census income."""
    out = []
    for r in D["cohort"]:
        inc = float(r["parents_total_income_2000"])
        if inc <= 0:
            continue
        b = int(r["birth_year"])
        e = D["earn"][r["child_id"]]
        vals = [e[b + a] * 100.0 / D["cpi"][b + a] for a in K.AGES]
        avg = sum(vals) / 5.0
        if avg <= 0:
            continue
        out.append(dict(child=r["child_id"], family=r["family_id"], region=D["recode"][r["census_region_2001"]],
                        y=np.log(avg), x=np.log(inc * 100.0 / D["cpi"][K.CENSUS_YEAR])))
    return out


def family_region(D):
    return {r["family_id"]: D["recode"][r["census_region_2001"]] for r in D["cohort"]}


def transitory(D):
    """Each region's variance of a family's log income around its own five-year level, over the region's families
    with a history (each family's squared deviations from its own mean, over four degrees of freedom a family)."""
    reg = family_region(D)
    num = {n: 0.0 for n in K.NAMES}
    den = {n: 0 for n in K.NAMES}
    for f, h in D["hist"].items():
        lx = np.log(np.array([h[y] for y in K.TAX_YEARS]))
        num[reg[f]] += float(((lx - lx.mean()) ** 2).sum())
        den[reg[f]] += len(lx) - 1
    return {n: num[n] / den[n] for n in K.NAMES}, {n: den[n] // (len(K.TAX_YEARS) - 1) for n in K.NAMES}


def elasticities(D):
    kids = children(D)
    s2, nhist = transitory(D)
    res = {}
    for name in K.NAMES:
        sub = [k for k in kids if k["region"] == name]
        y = np.array([k["y"] for k in sub])
        x = np.array([k["x"] for k in sub])
        v = float(np.var(x, ddof=1))
        c = float(np.cov(y, x)[0, 1])
        res[name] = dict(n=len(sub), beta=c / (v - s2[name]), ols=c / v, var=v, s2=s2[name], hist=nhist[name],
                         reliability=(v - s2[name]) / v)
    return res, s2, kids


# ---------------------------------------------------------------------------------------------- the asks
def offices(folder):
    return {r["office"]: r["region"] for r in read(folder, K.OFFICES) if r["region"]}


def _swap(d):
    y, m, dd = (int(v) for v in d.split("-"))
    return "%04d-%02d-%02d" % (y, dd, m) if dd <= 12 else None


def starts(folder):
    """Young people Bridges started, by the region of the office where each first enrolled: each record's
    predecessor is the record it continues (a transfer's transfer_from; a case carried into Northstar in April 2021
    continues the CaseTrack record left active at the same office with the same intake date and date of birth, the
    date of birth CaseTrack held sometimes with day and month swapped)."""
    rows = read(folder, K.PART)
    parent = {x["participant_ref"]: x["transfer_from"] for x in rows if x["transfer_from"]}
    left = [x for x in rows if x["source_system"] == "CaseTrack" and x["status"] == "Active"]
    carried = [x for x in rows if x["source_system"] == "Northstar" and x["intake_date"] < K.MIGRATION.isoformat()]
    for c in left:
        m = [n for n in carried if n["office"] == c["office"] and n["intake_date"] == c["intake_date"]
             and n["date_of_birth"] in (c["date_of_birth"], _swap(c["date_of_birth"]))]
        assert len(m) == 1, c["participant_ref"]
        parent[m[0]["participant_ref"]] = c["participant_ref"]
    office = {x["participant_ref"]: x["office"] for x in rows}
    roots = set()
    for x in rows:
        ref = x["participant_ref"]
        while ref in parent:
            ref = parent[ref]
        roots.add(ref)
    reg = offices(folder)
    out = {n: 0 for n in K.NAMES}
    for ref in roots:
        out[reg[office[ref]]] += 1
    return out


def spending(folder, st):
    """Spending in each region: its office's lines, the year-end accruals taken back (the central reversal each
    January undoes them, and the December invoices are posted in January), and the central team's costs (the central
    lines other than those reversals) shared by the young people each region started."""
    lines = read(folder, K.LEDGER)
    reg = offices(folder)
    office_tot = {o: 0.0 for o in reg}
    central = 0.0
    for l in lines:
        a = float(l["amount"])
        if l["cost_centre"] in reg:
            if l["description"] != "Year-end accrual, mentor fees":
                office_tot[l["cost_centre"]] += a
        elif not l["description"].startswith("Reversal of"):
            central += a
    tot = sum(st.values())
    return {name: office_tot[o] + central * st[name] / tot for o, name in reg.items()}


def thousands(x):
    return int(np.floor(x / 1000.0 + 0.5))


def compute(folder):
    D = load(folder)
    el, s2, kids = elasticities(D)
    order = sorted(K.NAMES, key=lambda n: -el[n]["beta"])
    st = starts(folder)
    sp = spending(folder, st)
    fig = {"call": "Opportunity Pathways goes to %s in 2027." % order[0], "region": order[0], "runner_up": order[1],
           "earnings_ages": "%d to %d" % (K.AGES[0], K.AGES[-1])}
    for n in K.NAMES:
        key = n.lower().replace(" ", "_")
        fig["elasticity_" + key] = "%.2f" % el[n]["beta"]
        fig["transitory_" + key] = "%.3f" % el[n]["s2"]
        fig["children_" + key] = "{:,}".format(el[n]["n"])
        fig["started_" + key] = "{:,}".format(st[n])
        fig["spending_" + key] = "${:,}".format(thousands(sp[n]) * 1000)
    fig["children_total"] = "{:,}".format(sum(el[n]["n"] for n in K.NAMES))
    fig["started_total"] = "{:,}".format(sum(st.values()))
    fig["spending_total"] = "${:,}".format(thousands(sum(sp.values())) * 1000)
    fig["gap"] = "%.2f" % (round(el[order[0]]["beta"], 2) - round(el[order[1]]["beta"], 2))
    fams = {k["family"] for k in kids}
    hist_region = {}
    reg = family_region(D)
    for n in K.NAMES:
        fn = {f for f in fams if reg[f] == n}
        hist_region[n] = sum(1 for f in fn if f in D["hist"]) / len(fn)
    hist_share = sum(1 for f in fams if f in D["hist"]) / len(fams)
    r0 = {}
    for n in K.NAMES:
        sub = [k for k in kids if k["region"] == n and k["family"] in D["hist"]]
        y = np.array([k["y"] for k in sub])
        xb = np.array([np.log(np.mean([D["hist"][k["family"]][t] for t in K.TAX_YEARS])) for k in sub])
        r0[n] = float(np.cov(y, xb)[0, 1]) / float(np.var(xb, ddof=1))
    return dict(el=el, s2=s2, order=order, starts=st, spending=sp, figures=fig, n_kids=len(kids),
                hist_region=hist_region, hist_share=hist_share, r0=r0, r0_top=max(r0, key=r0.get))


# ---------------------------------------------------------------------------------------------- the deliverables
def k_(x):
    return "$" + format(thousands(x) * 1000, ",")


def write_chart(path, res):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams["font.family"] = "DejaVu Sans"
    el, order = res["el"], res["order"]
    fig, ax = plt.subplots(figsize=(8.0, 4.6), dpi=110)
    ys = np.arange(len(order))[::-1]
    for y, n in zip(ys, order):
        v = round(el[n]["beta"], 2)
        col = "#1F5C4A" if n == order[0] else "#9DB3AA"
        ax.barh(y, v, height=0.62, color=col)
        ax.text(v + 0.006, y, "%.2f" % v, va="center", fontsize=9, color="#1D2329")
    ax.set_yticks(ys)
    ax.set_yticklabels(["%s (%s children)" % (n, format(el[n]["n"], ",")) for n in order], fontsize=9)
    ax.set_xlim(0, 0.6)
    ax.set_xlabel("Intergenerational elasticity of earnings at 30 to 34 on parents' long-run income", fontsize=8.8)
    fig.text(0.012, 0.965, "%s: children's earnings as adults depend most on their parents' long-run income" % order[0],
             fontsize=10.5, va="top", color="#1D2329")
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="x", color="#DDDDDD", lw=0.6)
    ax.set_axisbelow(True)
    fig.text(0.012, 0.015, "Every child born 1984 to 1988 in the region at the 2001 census. Parents' 2000 census income, its "
             "year-to-year swing (measured within\nfamilies over the 2001 to 2005 tax years) taken out of the spread. "
             "Source: %s, linked extract, 2026 release." % K.CENTRE, fontsize=7.0, color="#5C666F", va="bottom")
    fig.tight_layout(rect=(0, 0.07, 1, 0.94))
    fig.savefig(path, format="png", metadata={"Software": None})
    plt.close(fig)


def _numtable(DC, S, rows, widths, total=False):
    """The memo's tables: figures set right, a rule over a total row."""
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, Table, TableStyle
    right = ParagraphStyle("r", parent=S["cell"], alignment=2)
    rightb = ParagraphStyle("rb", parent=S["cellb"], alignment=2)
    data = []
    for i, r in enumerate(rows):
        bold = i == 0 or (total and i == len(rows) - 1)
        data.append([Paragraph(str(c), (S["cellb"] if bold else S["cell"]) if j == 0 else (rightb if bold else right))
                     for j, c in enumerate(r)])
    t = Table(data, colWidths=[w * DC.mm for w in widths], hAlign="LEFT")
    style = [("LINEBELOW", (0, 0), (-1, 0), 0.7, DC.INK), ("LINEBELOW", (0, 1), (-1, -1), 0.25, DC.RULE),
             ("VALIGN", (0, 0), (-1, -1), "BOTTOM"), ("TOPPADDING", (0, 0), (-1, -1), 2.6),
             ("BOTTOMPADDING", (0, 0), (-1, -1), 2.6)]
    if total:
        style.append(("LINEABOVE", (0, -1), (-1, -1), 0.7, DC.INK))
    t.setStyle(TableStyle(style))
    return t


def write_memo(path, res, chart):
    import docs as DC
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, KeepTogether, Paragraph, Spacer
    S = DC.styles("Sans", colors.HexColor("#1F5C4A"))
    el, order, st, sp = res["el"], res["order"], res["starts"], res["spending"]
    top, second = order[0], order[1]
    b0, b1 = round(el[top]["beta"], 2), round(el[second]["beta"], 2)
    story = [Paragraph("%s 2027: %s, where children's earnings depend most on their parents' income" % (
                 K.PROGRAM, top), S["h1"]),
             Paragraph("To: %s, Director of Programs<br/>From: Research and Evaluation<br/>Date: %s<br/>Subject: the "
                       "2027 region, for the board of trustees on %s" % (
                           K.PEOPLE["director"], dt.date(2026, 11, 18).strftime("%-d %B %Y"),
                           K.BOARD.strftime("%-d %B")), S["meta"]),
             Paragraph("Recommendation", S["h2"]),
             Paragraph("%s should go to %s. Of the province's eight regions, it is where children's earnings as adults "
                       "depend most on their parents' long-run income: an intergenerational elasticity of %.2f, against "
                       "%.2f in %s, the next highest, a gap of %.2f." % (K.PROGRAM, top, b0, b1, second, b0 - b1),
                       S["body"]),
             Paragraph("How the elasticity was measured", S["h2"]),
             Paragraph("The children are the %s born from 1984 to 1988 who lived with a parent in the province at the 2001 "
                       "census, had parents' census income above zero and had earnings in at least one year from 30 to "
                       "34, each placed in the region of the 2001 household. A child's earnings are the average of the "
                       "five years at 30 to 34 in 2005 dollars, years without earnings counted as zero." % format(
                           sum(el[n]["n"] for n in K.NAMES), ","), S["body"]),
             Paragraph("Long-run income is observed for no family. The census gives every family's income for one year, "
                       "2000, and one year's income moves around the family's long-run level: in each region part of the "
                       "spread of census incomes is that year-to-year swing, which flattens the slope of children's "
                       "earnings on it. The tax history measures the swing directly, as the variance of each family's "
                       "income around its own five-year level from 2001 to 2005, region by region. Taking that variance "
                       "out of the spread of the census incomes leaves the spread of long-run income, and each region's "
                       "elasticity is the covariance of children's earnings with the census income over that spread.",
                       S["body"]),
             Paragraph("The tax history itself covers only the families in which a parent filed in every year, %d per "
                       "cent of the children's families across the province and as few as %d per cent in %s. They are "
                       "the better-off families in most regions, so a slope measured on them alone describes a "
                       "narrower slice of the region and would put %s first. What carries from them to every family is "
                       "the size of the year-to-year swing, not how reliable one year's income is among them, since that "
                       "depends on how far apart their long-run incomes sit." % (
                           round(100 * res["hist_share"]), round(100 * min(res["hist_region"].values())),
                           min(res["hist_region"], key=res["hist_region"].get), res["r0_top"]), S["body"])]
    rows = [["Region", "Children", "Elasticity", "Census-year slope", "Year-to-year swing (variance of log income)"]]
    for n in order:
        rows.append([n, format(el[n]["n"], ","), "%.2f" % el[n]["beta"], "%.2f" % el[n]["ols"], "%.3f" % el[n]["s2"]])
    story.append(Spacer(1, 4))
    story.append(_numtable(DC, S, rows, [36, 20, 20, 28, 62]))
    story.append(Paragraph("Census-year slope: the slope on one year's income before the swing is taken out. Source: %s, "
                           "linked intergenerational income extract, 2026 release. Figures from the extract go to the "
                           "centre for its release check before the board papers are published (user guide, section 7)."
                           % K.CENTRE, S["note"]))
    story.append(KeepTogether([Spacer(1, 6), Image(chart, width=160 * mm, height=92 * mm)]))
    story.append(Paragraph("%s in each region, 2019 to 2023" % K.OLD_PROGRAM, S["h2"]))
    story.append(Paragraph(
        "Young people are counted once, in the region of the office where they first enrolled: a case carried from "
        "CaseTrack into Northstar in April 2021 and a move to another office continue the same young person. Spending "
        "is each office's ledger with the December accruals taken back (the central reversal each January undoes them "
        "and the December invoices are posted in January), plus the central team's costs shared by the young people "
        "each region started, through the close-out in March 2024.", S["body"]))
    rows = [["Region", "Young people started", "Spending"]]
    for n in K.NAMES:
        rows.append([n, format(st[n], ","), k_(sp[n])])
    rows.append(["Province", format(sum(st.values()), ","), k_(sum(sp.values()))])
    story.append(_numtable(DC, S, rows, [40, 36, 36], total=True))
    story.append(Paragraph("Source: %s participant export, program ledger and office list. Each region is rounded to the "
                           "thousand on its own; the province total is rounded from the unrounded sum." % K.OLD_PROGRAM,
                           S["note"]))
    doc = DC._doc(path, "%s 2027: %s" % (K.PROGRAM, top), "Research and Evaluation")
    doc.build(story, onFirstPage=DC._frame(K.FOUNDATION, "Research and Evaluation"),
              onLaterPages=DC._frame(K.FOUNDATION, "Research and Evaluation"))


def write_xlsx(path, res):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    el, order, st, sp = res["el"], res["order"], res["starts"], res["spending"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Regions"
    ws.append(["Region", "Children", "Intergenerational elasticity", "Census-year slope", "Variance of log census income",
               "Year-to-year swing (variance of log income)", "Families with a five-year history"])
    for n in order:
        ws.append([n, el[n]["n"], round(el[n]["beta"], 2), round(el[n]["ols"], 2), round(el[n]["var"], 3),
                   round(el[n]["s2"], 3), el[n]["hist"]])
    ws.append(["Province", sum(el[n]["n"] for n in K.NAMES), None, None, None, None, sum(el[n]["hist"] for n in K.NAMES)])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    s = wb.create_sheet("Bridges")
    s.append(["Region", "Young people started, 2019 to 2023", "Spending ($)"])
    for n in K.NAMES:
        s.append([n, st[n], thousands(sp[n]) * 1000])
    s.append(["Province", sum(st.values()), thousands(sum(sp.values())) * 1000])
    for c in s[s.max_row]:
        c.font = Font(bold=True)
    s = wb.create_sheet("Notes")
    notes = ["Recommendation: %s goes to %s in 2027 (elasticity %.2f; next, %s, %.2f)." % (
                 K.PROGRAM, order[0], round(el[order[0]]["beta"], 2), order[1], round(el[order[1]]["beta"], 2)),
             "Children: born 1984 to 1988, with a parent in the region at the 2001 census, parents' census income above "
             "zero, earnings in at least one year at 30 to 34 (the brief).",
             "Earnings: average of the five years at 30 to 34 in 2005 dollars, years without earnings as zero; logged.",
             "Parents' income: the 2000 census total, logged. Year-to-year swing: each region's variance of a family's log "
             "income around its own 2001 to 2005 level, over the families with a five-year history (four degrees of "
             "freedom a family).",
             "Elasticity: covariance of the child's log earnings with the log census income, over the variance of the log "
             "census income less the region's year-to-year swing.",
             "Bridges: young people counted once at the office of first enrolment (transfers and the April 2021 move to "
             "Northstar followed back); spending from the program ledger through March 2024, December accruals taken "
             "back, central team costs shared by young people started; rounded to the nearest thousand dollars.",
             "Regions' spending is rounded to the thousand on its own; the province total is rounded from the "
             "unrounded sum.",
             "Sources: %s linked extract (release of %s); %s participant export, ledger and office list." % (
                 K.CENTRE, K.RELEASE.strftime("%-d %B %Y"), K.OLD_PROGRAM),
             "Prepared by Research and Evaluation for %s, %s." % (K.PEOPLE["director"],
                                                                dt.date(2026, 11, 18).strftime("%-d %B %Y"))]
    for t in notes:
        s.append([t])
    s.column_dimensions["A"].width = 130
    for sh in wb.worksheets[:2]:
        for c in sh[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.fill = PatternFill("solid", fgColor="DCE9E3")
        sh.freeze_panes = "A2"
        for col in "ABCDEFG":
            sh.column_dimensions[col].width = 18
        sh.column_dimensions["A"].width = 20
        for row in sh.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, int) and c.value >= 1000:
                    c.number_format = "$#,##0" if (sh.title == "Bridges" and c.column == 3) else "#,##0"
                elif isinstance(c.value, float):
                    c.number_format = "0.00" if c.column in (3, 4) else "0.000"
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth = 1
        sh.sheet_properties.pageSetUpPr.fitToPage = True
        sh.print_title_rows = "1:1"
    wb.save(path)
