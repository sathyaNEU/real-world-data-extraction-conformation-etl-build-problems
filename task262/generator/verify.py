"""Independent verifier: reads only the shipped bundle (target/) and the goldens, with its own code, and recomputes the
call, every region's elasticity and child count, and the Bridges starts and spending, then compares them with the
golden workbook. No generator module is imported.

    python3 verify.py <target> <golden>"""
import csv
import math
import os
import sys
from collections import defaultdict

AGES = range(30, 35)


def rows(folder, name):
    with open(os.path.join(folder, name), encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def mean(v):
    return sum(v) / len(v)


def cov(a, b):
    ma, mb = mean(a), mean(b)
    return sum((x - ma) * (y - mb) for x, y in zip(a, b)) / (len(a) - 1)


def elasticities(t):
    cpi = {int(r["year"]): float(r["cpi_2005_100"]) for r in rows(t, "consumer_price_index_2000_2022.csv")}
    name = {r["census_region_2001"]: r["region_2021"] for r in rows(t, "region_recode_2001_2021.csv")}
    earn = defaultdict(dict)
    for r in rows(t, "child_earnings_2014_2022.csv"):
        earn[r["child_id"]][int(r["tax_year"])] = int(r["earnings"])
    fam_region = {}
    xs, ys = defaultdict(list), defaultdict(list)
    for r in rows(t, "linked_cohort_extract_2026.csv"):
        reg = name[r["census_region_2001"]]
        fam_region[r["family_id"]] = reg
        inc = int(r["parents_total_income_2000"])
        if inc <= 0:
            continue
        b = int(r["birth_year"])
        avg = sum(earn[r["child_id"]][b + a] * 100.0 / cpi[b + a] for a in AGES) / 5
        if avg <= 0:
            continue
        xs[reg].append(math.log(inc * 100.0 / cpi[2000]))
        ys[reg].append(math.log(avg))
    hist = defaultdict(list)
    for r in rows(t, "parent_income_history_2001_2005.csv"):
        hist[r["family_id"]].append(math.log(int(r["family_total_income_2005_dollars"])))
    ss, df = defaultdict(float), defaultdict(int)
    for f, v in hist.items():
        assert len(v) == 5, f
        m = mean(v)
        ss[fam_region[f]] += sum((a - m) ** 2 for a in v)
        df[fam_region[f]] += 4
    out = {}
    for reg in xs:
        s2 = ss[reg] / df[reg]
        out[reg] = (cov(ys[reg], xs[reg]) / (cov(xs[reg], xs[reg]) - s2), len(xs[reg]))
    return out


def starts(t):
    recs = rows(t, "bridges_participants_2019_2023.csv")
    region = {r["office"]: r["region"] for r in rows(t, "bridges_offices.csv") if r["region"]}
    prev = {r["participant_ref"]: r["transfer_from"] for r in recs if r["transfer_from"]}

    def swapped(d):
        y, m, dd = d.split("-")
        return "-".join((y, dd, m)) if int(dd) <= 12 else None

    old = [r for r in recs if r["source_system"] == "CaseTrack" and r["status"] == "Active"]
    new = [r for r in recs if r["source_system"] == "Northstar" and r["intake_date"] < "2021-04-06"]
    for o in old:
        hits = [n for n in new if (n["office"], n["intake_date"]) == (o["office"], o["intake_date"])
                and n["date_of_birth"] in (o["date_of_birth"], swapped(o["date_of_birth"]))]
        assert len(hits) == 1, o["participant_ref"]
        prev[hits[0]["participant_ref"]] = o["participant_ref"]
    office = {r["participant_ref"]: r["office"] for r in recs}
    first = set()
    for r in recs:
        ref = r["participant_ref"]
        while ref in prev:
            ref = prev[ref]
        first.add(ref)
    out = defaultdict(int)
    for ref in first:
        out[region[office[ref]]] += 1
    return dict(out)


def spending(t, st):
    region = {r["office"]: r["region"] for r in rows(t, "bridges_offices.csv") if r["region"]}
    office = defaultdict(float)
    team = 0.0
    for r in rows(t, "bridges_program_ledger_2019_2024.csv"):
        a = float(r["amount"])
        if r["cost_centre"] in region:
            if not r["description"].startswith("Year-end accrual"):
                office[region[r["cost_centre"]]] += a
        elif not r["description"].startswith("Reversal of"):
            team += a
    n = sum(st.values())
    return {reg: office[reg] + team * st[reg] / n for reg in region.values()}


def main(target, golden):
    el = elasticities(target)
    order = sorted(el, key=lambda k: -el[k][0])
    st = starts(target)
    sp = spending(target, st)
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(golden, "opportunity_pathways_2027.xlsx"), data_only=True)
    reg_rows = {r[0]: r for r in wb["Regions"].iter_rows(min_row=2, values_only=True)}
    br_rows = {r[0]: r for r in wb["Bridges"].iter_rows(min_row=2, values_only=True)}
    bad = []
    for k in el:
        if reg_rows[k][1] != el[k][1] or abs(reg_rows[k][2] - round(el[k][0], 2)) > 1e-9:
            bad.append(("elasticity", k))
        if br_rows[k][1] != st[k] or br_rows[k][2] != int(math.floor(sp[k] / 1000 + 0.5)) * 1000:
            bad.append(("bridges", k))
    first_row = [r for r in wb["Regions"].iter_rows(min_row=2, values_only=True)][0][0]
    if first_row != order[0]:
        bad.append(("call", order[0], first_row))
    if bad:
        raise SystemExit("verify: mismatches %s" % bad)
    print("verify: the call (%s, %.2f; next %s, %.2f), 8 elasticities and child counts, 8 start counts and 8 spending "
          "figures recomputed from the bundle alone and equal to the golden workbook" % (
              order[0], el[order[0]][0], order[1], el[order[1]][0]))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
