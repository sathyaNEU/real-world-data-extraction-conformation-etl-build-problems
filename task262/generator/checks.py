"""The generator's assertions on every build: the three rules (R1 raw material, R2 each mess item neutral, R3a the
analyst test, R3b the decisive step alone, R3c no plant), the ladder and its reflexes, every row of the semantic fork
table (SF1 to SF12), the estimators' convergence and every graded figure's distance to its rounding edge, the asks'
goldens against the generator's truth and their cracker sheet, the separation line, the distractors, the input gates
and the realism sweep."""
from __future__ import annotations

import collections
import csv
import os
import re
import shutil
import tempfile
import zipfile

import numpy as np

import bridges as BR
import common as K
import deliver as G
import stops as SP
import world as W

B = "Selvane Valley"
PASSED = []


def check(label, cond, detail=""):
    if not cond:
        raise AssertionError("%s %s" % (label, detail))
    PASSED.append(label)


def ranks(vals, key):
    return sorted(K.NAMES, key=lambda n: -vals[n][key])


def edge(v, unit=0.01):
    """Distance from v to the nearest rounding edge at this unit."""
    f = v / unit - np.floor(v / unit)
    return abs(f - 0.5) * unit


def run(target, golden, st, meta_path):
    del PASSED[:]
    S = st["S"]
    pk = SP.Pack(target)
    res = st["res"]
    el = res["el"]
    R, pool = SP.rungs(pk)

    # ------------------------------------------------------------------------------------------ R: the three rules
    cols = {}
    for f in (K.COHORT, K.EARN, K.HIST, K.RECODE, K.CPI_FILE, K.PART, K.LEDGER, K.OFFICES):
        with open(os.path.join(target, f), encoding="utf-8") as fh:
            cols[f] = next(csv.reader(fh))
    check("R1 the field's raw material ships: one row per child with the parents' census income, one row per child and "
          "year of earnings, one row per family and tax year of income",
          cols[K.COHORT][-1] == "parents_total_income_2000" and cols[K.EARN] == ["child_id", "tax_year", "earnings"]
          and cols[K.HIST] == ["family_id", "tax_year", "family_total_income_2005_dollars"])
    banned = re.compile(r"elastic|slope|varian|reliab|permanent|long.?run|transitory|rank", re.I)
    check("R1 no shipped column holds the deciding quantity (no elasticity, slope, variance, reliability or long-run "
          "income column)", not any(banned.search(c) for f in cols for c in cols[f]))
    check("R1 the golden's decisive computation reads the history's rows: every region's transitory variance comes from "
          "its families' five tax years", all(pk.nhist[n] > 500 for n in K.NAMES) and
          sum(pk.nhist.values()) * 5 == st["counts"][K.HIST])
    # R2: each Layer 1 item mishandled leaves the call where it was
    conv = SP.conventions(pk)
    for key, label in (("SF8", "census income left in 2000 dollars and earnings in current dollars (no deflation)"),
                       ("SF3", "years without earnings dropped from the average"),
                       ("SF7", "one child per family (siblings read as repeats)")):
        o = sorted(K.NAMES, key=lambda n: -conv[key][n])
        check("R2 mess neutral: %s keeps %s" % (label, B), o[0] == B and conv[key][o[0]] / conv[key][o[1]] > 1.15,
              str(o[:2]))
    fixed = fixed_window(pk)
    o = sorted(K.NAMES, key=lambda n: -fixed[n])
    check("R2 mess neutral: one calendar window (2018 to 2022) for every cohort instead of the ages 30 to 34 keeps %s" % B,
          o[0] == B and fixed[o[0]] / fixed[o[1]] > 1.15)
    recode = {r["census_region_2001"]: r["region_2021"] for r in G.read(target, K.RECODE)}
    check("R2 mess neutral: the 2001 codes and the 2021 names group the same children (one to one, no boundary change)",
          len(set(recode.values())) == len(recode) == 8 and set(recode.values()) == set(K.NAMES))
    nonpos = sum(1 for r in pk.D["cohort"] if float(r["parents_total_income_2000"]) <= 0)
    check("R2 mess neutral: census incomes at or below zero (business and farm losses) are %d rows, under one per cent, "
          "and the brief leaves them out" % nonpos, 0 < nonpos < 0.01 * len(pk.D["cohort"]))
    clean, _ = W.routes(S["regions"])
    o0 = ranks(clean, "r0")
    check("R3a the analyst test on the clean truth: the filers' five-year averages regressed on by region file Halverston",
          o0[0] == "Halverston" and clean[o0[0]]["r0"] / clean[o0[1]]["r0"] > 1.12, str(o0[:2]))
    og = ranks(clean, "gold")
    check("R3b the decisive step alone on the clean truth (every child, the census year less the region's transitory "
          "variance) files %s" % B, og[0] == B and clean[og[0]]["gold"] / clean[og[1]]["gold"] > 1.15)
    for key, lab in (("r1", "every child on the census year, nothing taken out"),
                     ("r2", "the census year over the filers' reliability ratio"),
                     ("r1c", "the filers' averages with their own residual swing taken out")):
        o = ranks(clean, key)
        check("R3b one correction alone stays wrong: %s files %s, not %s" % (lab, o[0], B), o[0] != B)
    hist = G.read(target, K.HIST)
    fams = collections.defaultdict(list)
    for r in hist:
        fams[r["family_id"]].append(r["family_total_income_2005_dollars"])
    blocks = collections.Counter(tuple(v) for v in fams.values())
    check("R3c no plant: no family's five-year income history repeats another's", blocks.most_common(1)[0][1] == 1)
    seqs = collections.Counter()
    for c, e in pk.D["earn"].items():
        if any(v > 0 for v in e.values()):
            seqs[tuple(e[y] for y in K.EARN_YEARS)] += 1
    check("R3c no plant: no child's nine years of earnings repeat another child's (all-zero rows aside)",
          seqs.most_common(1)[0][1] == 1)
    texts = all_text(target)
    check("R3c no plant: no shipped file publishes an elasticity, a reliability ratio, a transitory variance or a ranking "
          "of the regions on persistence", not re.search(r"elasticit[y-z]* (of|=|is) ?0\.|reliability ratio|transitory|"
                                                          r"persistence", " ".join(v for k, v in texts.items() if k != K.BRIEF), re.I))
    check("R3c no plant: the brief names no region and no candidate figure",
          not any(n in texts[K.BRIEF] for n in K.NAMES) and not re.search(r"\b0\.\d", texts[K.BRIEF]))

    # ------------------------------------------------------------------------------------------ the ladder
    lad = {}
    for key, win, mm in (("r0", "Halverston", 1.15), ("r1", "Torvale Metro", 1.15), ("r2", "Kestry North", 1.15),
                         ("gold", B, 1.2)):
        o, m = SP.winner(R, key)
        lad[key] = (o, m)
        check("ladder %s names %s by %.3f (floor %.2f)" % (key, win, m, mm), o[0] == win and m >= mm, str(o[:2]))
    check("ladder: every rung names a different region", len({lad[k][0][0] for k in lad}) == 4)
    check("position: %s is 4th or lower on R0 and R1 and never first before R3" % B,
          lad["r0"][0].index(B) >= 3 and lad["r1"][0].index(B) >= 3 and lad["r2"][0].index(B) >= 1)
    check("position: %s second on no rung" % B, all(lad[k][0].index(B) != 1 for k in ("r0", "r1", "r2")))
    gold = {n: R[n]["gold"] for n in K.NAMES}
    check("the golden on the pack equals deliver's figures to the digit", all(
        abs(gold[n] - el[n]["beta"]) < 1e-12 for n in K.NAMES))
    mult = {n: R[n]["gold"] / R[n]["r1"] for n in K.NAMES}
    w1 = lad["r1"][0][0]
    lead = R[w1]["r1"] / R[B]["r1"]
    check("dominance: the decisive move raises %s's census slope %.2f times against %s's %.2f, an edge of %.2f, at least "
          "1.2 times %s's census-year lead of %.2f" % (B, mult[B], w1, mult[w1], mult[B] / mult[w1], w1, lead),
          mult[B] / mult[w1] >= 1.2 * lead)
    o, m = SP.winner(R, "pool")
    check("reflex: a provincial transitory variance keeps the call on %s (%.3f)" % (B, m), o[0] == B and m > 1.1)
    for key, lab in (("r1c", "the filers' averages with their residual swing removed"),
                     ("iv", "the filers' averages instrumented by the census year"),
                     ("ipw", "the filers' averages reweighted to every family's chance of filing"),
                     ("rank", "the rank-rank slope")):
        o = ranks(R, key)
        check("reflex: %s names %s, %s at %.2f of it" % (lab, o[0], B, R[B][key] / R[o[0]][key]),
              o[0] != B and R[o[0]][key] / R[B][key] > 1.08)

    # ------------------------------------------------------------------------------------------ the fork table
    check("SF1 the region: the 2001 code through the recode, one to one; the extract carries no later address",
          "census_region_2001" in cols[K.COHORT] and not any("2022" in c or "address" in c for c in cols[K.COHORT]))
    o = sorted(K.NAMES, key=lambda n: -conv["SF2"][n])
    check("SF2 earnings in all five years names %s, ruled out by the brief (earnings in at least one year, a year without "
          "earnings counting as zero)" % o[0], o[0] != B and "at least one year" in texts[K.BRIEF]
          and "counting as zero" in texts[K.BRIEF])
    o = sorted(K.NAMES, key=lambda n: -conv["SF3"][n])
    check("SF3 the average over earning years only: same call", o[0] == B)
    check("SF4 the five-year average read as long-run income is the default path (Halverston), ruled out by the brief's "
          "every child and its long-run level", lad["r0"][0][0] == "Halverston" and "every child" in texts[K.BRIEF]
          and "the level its yearly income moves around" in texts[K.BRIEF])
    check("SF5 parents' income: the extract carries the parents' total only", "parents_total_income_2000" in cols[K.COHORT]
          and not any("household_income" in c for c in cols[K.COHORT]))
    o = ranks(R, "rank")
    check("SF6 the rank-rank slope names %s, ruled out by the brief's elasticity (the slope of the logarithm)" % o[0],
          o[0] != B and "logarithm" in texts[K.BRIEF] and "elasticity" in texts[K.BRIEF])
    o = sorted(K.NAMES, key=lambda n: -conv["SF7"][n])
    check("SF7 one child per family: same call", o[0] == B)
    o = sorted(K.NAMES, key=lambda n: -conv["SF8"][n])
    check("SF8 nominal dollars: same call, same two-decimal elasticities",
          o[0] == B and all(round(conv["SF8"][n], 2) == round(gold[n], 2) for n in K.NAMES))
    sd = {n: pk.s2[n] for n in K.NAMES}
    check("SF9 the transitory variance's scope: a provincial variance gives the same call; the regions' own differ from "
          "%.3f to %.3f on %d to %d families each, so each region's is its own" % (
              min(sd.values()), max(sd.values()), min(pk.nhist.values()), max(pk.nhist.values())),
          SP.winner(R, "pool")[0][0] == B and max(sd.values()) / min(sd.values()) > 1.8 and min(pk.nhist.values()) > 1000)
    check("SF10 the number of children: the children each elasticity rests on, summing to the brief's population",
          sum(el[n]["n"] for n in K.NAMES) == res["n_kids"] < st["counts"][K.COHORT])
    E = SP.estimators(pk)
    check("SF11 rounding: every elasticity at least 0.0015 from a two-decimal edge under every estimator", all(
        edge(v) >= 0.0015 for n in K.NAMES for v in E[n]), str({n: min(edge(v) for v in E[n]) for n in K.NAMES}))
    check("SF12 the census year as a sixth history year: same two-decimal elasticities",
          all(round(E[n][2], 2) == round(E[n][0], 2) for n in K.NAMES))
    check("C1 the transitory variance by within-family deviations or by half the squared year-to-year changes gives the "
          "same two-decimal elasticities", all(round(E[n][1], 2) == round(E[n][0], 2) for n in K.NAMES))
    check("C1 the transitory variance with each family weighted by its children in the sample, or over the families with "
          "a child in the sample only, gives the same two-decimal elasticities",
          all(round(E[n][i], 2) == round(E[n][0], 2) for n in K.NAMES for i in (3, 4)))
    order = res["order"]
    check("runner-up: %s, above the third by %.3f or more under every estimator" % (order[1], min(
        E[order[1]][i] - max(E[n][i] for n in order[2:]) for i in range(5))), all(
        sorted(K.NAMES, key=lambda n: -E[n][i])[:2] == order[:2] for i in range(5)) and min(
        E[order[1]][i] - max(E[n][i] for n in order[2:]) for i in range(5)) >= 0.01)

    # ------------------------------------------------------------------------------------------ the asks
    truth_st = {BR.OFFICE_REGION[o]: v for o, v in BR.truth_starts(S["P"]).items()}
    check("Ask A golden equals the generator's young people by first office", res["starts"] == truth_st)
    truth_sp = BR.truth_spending(S["lines"], S["P"])
    check("Ask B golden equals the generator's spending by region to the cent",
          all(abs(res["spending"][n] - truth_sp[n]) < 0.01 for n in K.NAMES))
    A = SP.ask_a_stops(target)
    for key, lab in (("natural", "every record counted"), ("chain", "transfer chains followed (the guidelines)"),
                     ("dedup", "transfer chains followed, rows repeating office, intake date and date of birth dropped "
                               "(the cracker)")):
        check("Ask A stop misses every region: %s" % lab, all(A[key][n] != truth_st[n] for n in K.NAMES),
              str({n: A[key][n] - truth_st[n] for n in K.NAMES}))
    swapped = collections.Counter(BR.OFFICE_REGION[p["office"]] for p in S["P"] if p.get("swapped"))
    check("Ask A device: every region has a carried case whose CaseTrack date of birth has day and month swapped",
          all(swapped[n] >= 2 for n in K.NAMES), str(dict(swapped)))
    check("Ask A device is fair: each CaseTrack record left active pairs with exactly one Northstar record (asserted "
          "inside deliver.starts)", sum(1 for x in S["rows"] if x["system"] == "CaseTrack" and x["status"] == "Active")
          == sum(1 for p in S["P"] if p.get("migrated")))
    Bst, info = SP.ask_b_stops(target, {"golden": truth_st, "dedup": A["dedup"], "chain": A["chain"],
                                        "dropped": A["dropped"]})
    for key in ("shared reversal, starts golden", "shared reversal, starts dedup", "shared reversal, starts chain",
                "shared reversal, starts dropped", "negative lines dropped", "2019 to 2023 postings", "office lines alone"):
        errs = [abs(Bst[key][n] - truth_sp[n]) for n in K.NAMES]
        check("Ask B stop misses every region by %s or more: %s" % (format(int(min(errs)), ","), key), min(errs) >= 1500)
    check("Ask B device: the central reversals equal the offices' year-end accruals to the cent",
          abs(info["reversals"] + sum(info["accruals"].values())) < 0.01)
    check("Ask B device: every office accrued in at least one year and every year had accruals in two or more offices",
          all(info["accruals"][n] > 0 for n in K.NAMES) and min(
              sum(1 for o in BR.accruing() if y in BR.accruing()[o]) for y in (2019, 2020, 2021, 2022)) >= 2)
    check("Ask B rounding: every region's spending at least $150 from a thousand-dollar edge",
          all(edge(res["spending"][n], 1000) >= 150 for n in K.NAMES))
    check("Ask golden recomputes with every call file removed, and the call with every ask file removed", separation(target))

    # ------------------------------------------------------------------------------------------ distractors, gates
    import json
    meta = json.load(open(meta_path))
    check("distractors: two or more named in metadata.json, each in the pack", len(meta["distractor_files"]) >= 2 and all(
        os.path.exists(os.path.join(target, f)) for f in meta["distractor_files"]))
    for f in meta["distractor_files"]:
        check("distractor deletion test: %s deleted, the golden and every ask unchanged" % f, without(target, f, res))
    files = sorted(os.listdir(target))
    exts = {f.rsplit(".", 1)[1] for f in files}
    check("input gates: %d files, %d formats, a file of %s rows" % (len(files), len(exts), format(st["counts"][K.EARN], ",")),
          len(files) >= 10 and len(exts) >= 3 and max(st["counts"].values()) >= 25000)
    check("input gates: no file name carries the stump or a distractor label", not any(
        re.search(r"distract|decoy|known_issue|answer|transitory|elastic", f) for f in files))

    # ------------------------------------------------------------------------------------------ realism sweep
    counts = [el[n]["n"] for n in K.NAMES]
    check("realism: no two regions share a child count, a start count or a history count",
          len(set(counts)) == 8 and len(set(res["starts"].values())) == 8 and len(set(pk.nhist.values())) == 8)
    totals = [sum(counts), sum(res["starts"].values()), st["counts"][K.COHORT], st["counts"][K.HIST] // 5,
              int(sum(res["spending"].values()))]
    check("realism: no round headline totals", all(t % 100 for t in totals), str(totals))
    check("realism: personas are the drawn ones", all(
        p in texts[K.BRIEF] + texts[K.GUIDE] + texts[K.DICT] for p in (K.PEOPLE["director"], K.PEOPLE["chair"])) and
        K.PEOPLE["analyst"] in texts[K.GUIDE] and not re.search(r"John Doe|Jane Doe|Acme|Company A", " ".join(texts.values())))
    leaks = [f for f in files + [os.path.join("..", "golden", g) for g in os.listdir(golden)]
             if re.search(rb"ReportLab|reportlab|openpyxl|python-docx|matplotlib|Matplotlib",
                          raw(os.path.join(target, f)))]
    check("realism: containers scrubbed of the writing libraries", not leaks, str(leaks))
    mt = {os.path.getmtime(os.path.join(target, f)) for f in files}
    check("realism: one in-fiction export time on every file", len(mt) == 1)
    check("realism: the user guide's counts match the files", format(st["counts"][K.COHORT], ",") in texts[K.GUIDE]
          and format(st["counts"][K.HIST] // 5, ",") in texts[K.GUIDE])
    by_year = st["closeout"][0]
    ledger_year = collections.defaultdict(float)
    for l in G.read(target, K.LEDGER):
        ledger_year[int(l["posting_date"][:4])] += float(l["amount"])
    check("context artifact: the close-out's spending by year is the ledger as posted, every cost centre",
          all(abs(by_year[y] - ledger_year[y]) < 0.01 for y in ledger_year))
    check("context artifact: the close-out's total ties under every handling of the asks (accruals and reversals net to "
          "zero over the program)", abs(sum(by_year.values()) - sum(res["spending"].values())) < 0.01)

    # ------------------------------------------------------------------------------------------ figures
    fresh = G.compute(target)
    check("figures.json recomputes from the pack", fresh["figures"] == res["figures"])
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(golden, K.DELIVERABLES[1]), data_only=True)
    ws = wb["Regions"]
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] in K.NAMES}
    check("golden workbook: each region's children and elasticity match figures.json", all(
        rows[n][1] == el[n]["n"] and "%.2f" % rows[n][2] == res["figures"]["elasticity_" + n.lower().replace(" ", "_")]
        for n in K.NAMES))
    ws = wb["Bridges"]
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] in K.NAMES}
    check("golden workbook: each region's starts and spending match figures.json", all(
        rows[n][1] == res["starts"][n] and "${:,}".format(rows[n][2]) == res["figures"]["spending_" + n.lower().replace(" ", "_")]
        for n in K.NAMES))
    check("H21 invented names clear of real places and organisations (no name within one edit of a real one, none "
          "containing one)", h21(), "")
    print("checks: %d assertions passed" % len(PASSED))
    return PASSED


# real places and organisations the invented names were screened against (the nearest places in a gazetteer of
# 5,600 British and Irish places, Canadian cities and regions, and the bodies the fiction sits beside)
REAL = ["Alveston", "Astley", "Broughton Astley", "Calverton", "Denholm", "Eastry", "Halberton", "Haswell", "Kelty",
        "Kerry", "Mosterton", "Otley", "Otterton", "Raskelf", "Ulverston", "Wedmore", "Wycombe", "High Wycombe",
        "Harlow", "Marlow", "Barlow", "Fenwick", "Ramsey", "Romsey", "Corran", "Corris", "North Shields", "Toronto",
        "Ottawa", "Montreal", "Vancouver", "Calgary", "Edmonton", "Winnipeg", "Regina", "Saskatoon", "Halifax",
        "Victoria", "Kelowna", "Kamloops", "Sudbury", "Moncton", "Fredericton", "Kingston", "Peterborough", "Guelph",
        "Hamilton", "Windsor", "Sarnia", "Nanaimo", "Whitehorse", "Yellowknife", "Revelstoke", "Lamont", "Interlake",
        "Parkland", "Kildonan", "Selkirk", "Muskoka", "Haliburton", "Algoma", "Kenora", "Timmins", "Cochrane",
        "Okanagan", "Kootenay", "Cariboo", "Skeena", "Athabasca", "Lloydminster", "Yorkton", "Dauphin", "Steinbach",
        "Statistics Canada", "Canadian Research Data Centre Network", "Maytree Foundation", "McConnell Foundation",
        "Laidlaw Foundation", "Metcalf Foundation", "Muttart Foundation", "Lawson Foundation", "Trillium Foundation",
        "Max Bell Foundation", "CaseWorthy", "Apricot", "Penelope"]
GENERIC = {"metro", "north", "valley", "lakes", "coast", "plains", "port", "falls", "harbour", "foundation", "inlet",
           "mills", "prairie", "head", "office", "research", "data", "centre", "provincial", "opportunity", "pathways"}


def h21():
    def lev(a, b):
        prev = list(range(len(b) + 1))
        for i, ca in enumerate(a, 1):
            cur = [i]
            for j, cb in enumerate(b, 1):
                cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
            prev = cur
        return prev[-1]
    import bridges as BRx
    invented = list(K.NAMES) + [t for _, _, t in BRx.OFFICES] + [K.FOUNDATION, K.CENTRE, K.PROGRAM, "Kestry-Inlet",
                                                                  "Selvane-Mills", "Ostler-Prairie"]
    words = {w for n in invented for w in re.sub(r"[^a-z ]", " ", n.lower()).split() if w not in GENERIC}
    real = set()
    for r in REAL:
        rn = re.sub(r"[^a-z ]", " ", r.lower())
        real.add(rn.replace(" ", ""))
        real.update(x for x in rn.split() if len(x) >= 4 and x not in GENERIC)
    for w in words:
        for r in real:
            if lev(w, r) <= 1 or (len(r) >= 5 and (r in w or w in r)):
                return False
    return True


def fixed_window(pk):
    D = pk.D
    out = {}
    kids = []
    for r in D["cohort"]:
        inc = float(r["parents_total_income_2000"])
        if inc <= 0:
            continue
        e = D["earn"][r["child_id"]]
        avg = sum(e[y] * 100.0 / D["cpi"][y] for y in range(2018, 2023)) / 5.0
        if avg <= 0:
            continue
        kids.append((D["recode"][r["census_region_2001"]], np.log(avg), np.log(inc * 100.0 / D["cpi"][K.CENSUS_YEAR])))
    for n in K.NAMES:
        y = np.array([k[1] for k in kids if k[0] == n])
        x = np.array([k[2] for k in kids if k[0] == n])
        out[n] = SP.cov(y, x) / (SP.var(x) - pk.s2[n])
    return out


def raw(path):
    return open(path, "rb").read()


def all_text(folder):
    """Every document's text: PDFs by pdftotext, OOXML by their XML parts with tags stripped, CSVs as they are."""
    import subprocess
    out = {}
    for f in os.listdir(folder):
        p = os.path.join(folder, f)
        if f.endswith(".pdf"):
            out[f] = subprocess.run(["pdftotext", "-layout", p, "-"], capture_output=True, text=True).stdout
            out[f] = re.sub(r"\s+", " ", out[f])
        elif f.endswith((".docx", ".xlsx")):
            with zipfile.ZipFile(p) as z:
                xml = " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml"))
            out[f] = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", xml))
        else:
            out[f] = open(p, encoding="utf-8").read(200000)
    return out


def without(target, name, res):
    tmp = tempfile.mkdtemp()
    try:
        for f in os.listdir(target):
            if f != name:
                shutil.copy(os.path.join(target, f), tmp)
        r2 = G.compute(tmp)
        return r2["figures"] == res["figures"]
    finally:
        shutil.rmtree(tmp)


def separation(target):
    tmp = tempfile.mkdtemp()
    try:
        call = (K.COHORT, K.EARN, K.HIST, K.RECODE, K.CPI_FILE)
        for f in os.listdir(target):
            if f not in call:
                continue
            shutil.copy(os.path.join(target, f), tmp)
        D = G.load(tmp)
        el, _, _ = G.elasticities(D)
        ref, _, _ = G.elasticities(G.load(target))
        same_call = all(abs(el[n]["beta"] - ref[n]["beta"]) < 1e-12 for n in K.NAMES)
        shutil.rmtree(tmp)
        os.makedirs(tmp)
        for f in (K.PART, K.LEDGER, K.OFFICES):
            shutil.copy(os.path.join(target, f), tmp)
        st = G.starts(tmp)
        return same_call and st == G.starts(target) and G.spending(tmp, st) == G.spending(target, st)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
