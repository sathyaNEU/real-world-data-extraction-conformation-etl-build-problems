"""The pack's documents: the Provincial Research Data Centre's user guide for the linked
extract, the director's brief to the trustees, the Bridges guidelines (2019) and close-out summary (2024), and two
papers the folder also holds (the 2021 regional economic indicators, the 2023 youth aspirations survey)."""
from __future__ import annotations

import datetime as dt

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                   # noqa: E402
from reportlab.lib.pagesizes import letter                                         # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                    # noqa: E402
from reportlab.lib.units import mm                                                 # noqa: E402
from reportlab.pdfbase import pdfmetrics                                           # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                       # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import common as K   # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("Sans", "LiberationSans-Regular"), ("Sans-B", "LiberationSans-Bold"), ("Sans-I", "LiberationSans-Italic")):
    pdfmetrics.registerFont(TTFont(nm, FONT_DIR + f + ".ttf"))
pdfmetrics.registerFontFamily("Sans", normal="Sans", bold="Sans-B", italic="Sans-I", boldItalic="Sans-B")
for nm, f in (("Serif", "LiberationSerif-Regular"), ("Serif-B", "LiberationSerif-Bold"), ("Serif-I", "LiberationSerif-Italic")):
    pdfmetrics.registerFont(TTFont(nm, FONT_DIR + f + ".ttf"))
pdfmetrics.registerFontFamily("Serif", normal="Serif", bold="Serif-B", italic="Serif-I", boldItalic="Serif-B")
INK, GREY, RULE = colors.HexColor("#1d2329"), colors.HexColor("#5c666f"), colors.HexColor("#b4bdc5")


def styles(font, accent):
    body = ParagraphStyle("b", fontName=font, fontSize=9.6, leading=13.4, textColor=INK, spaceAfter=5)
    return dict(
        body=body,
        h1=ParagraphStyle("h1", fontName=font + "-B", fontSize=15, leading=19, textColor=accent, spaceAfter=2),
        h2=ParagraphStyle("h2", fontName=font + "-B", fontSize=10.5, leading=14, textColor=INK, spaceBefore=8, spaceAfter=3),
        meta=ParagraphStyle("m", fontName=font, fontSize=8.3, leading=11, textColor=GREY, spaceAfter=9),
        note=ParagraphStyle("n", fontName=font, fontSize=7.8, leading=10, textColor=GREY, spaceBefore=6),
        cell=ParagraphStyle("c", fontName=font, fontSize=8.6, leading=11.2, textColor=INK),
        cellb=ParagraphStyle("cb", fontName=font + "-B", fontSize=8.6, leading=11.2, textColor=INK))


def _doc(path, title, author):
    return SimpleDocTemplate(path, pagesize=letter, leftMargin=22 * mm, rightMargin=22 * mm, topMargin=22 * mm,
                             bottomMargin=19 * mm, title=title, author=author, subject="", creator="", producer="")


def _frame(left, right, font="Sans"):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(GREY)
        canvas.setFont(font, 7.4)
        canvas.drawString(22 * mm, letter[1] - 12 * mm, left)
        canvas.drawRightString(letter[0] - 22 * mm, letter[1] - 12 * mm, right)
        canvas.drawRightString(letter[0] - 22 * mm, 10.5 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _table(S, rows, widths):
    data = [[Paragraph(str(c), S["cellb"] if i == 0 else S["cell"]) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=[w * mm for w in widths], hAlign="LEFT")
    t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, 0), 0.7, INK), ("LINEBELOW", (0, 1), (-1, -1), 0.25, RULE),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 2.6),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 2.6)]))
    return t


# ---------------------------------------------------------------------------------------------- the data centre
def write_user_guide(path, n_children, n_families, n_history):
    S = styles("Serif", colors.HexColor("#27455c"))
    story = [Paragraph("Linked intergenerational income extract, 2026 release: user guide", S["h1"]),
             Paragraph("%s. Release of %s. Contact: %s, research analyst." % (
                 K.CENTRE, K.RELEASE.strftime("%-d %B %Y"), K.PEOPLE["analyst"]), S["meta"])]
    sec = [
        ("1. Who is in the extract", [
            "The extract follows the %s children born from 1984 to 1988 who lived with at least one parent in a "
            "household that answered the long-form questionnaire of the 2001 Census of Population in the province. "
            "The long-form questionnaire went to one household in five, drawn at the same rate everywhere, so each "
            "child in the extract stands for the same number of children and no weights are needed. The %s families "
            "are the households the children lived in at the census." % (format(n_children, ","), format(n_families, ",")),
            "Each child is placed in the census region of the 2001 household. The census regions were renamed in 2021; "
            "the recode file gives each 2001 code its 2021 name. No boundary changed."]),
        ("2. Parents' income in the 2001 census", [
            "The census asked each adult's total income in 2000, from all sources and before tax, in the categories "
            "of income tax total income. The extract carries the total for the child's parents in the household "
            "(one or two parents), in 2000 dollars. Total income can be zero or below where a business or farm lost "
            "money."]),
        ("3. Parents' income history, 2001 to 2005", [
            "For families in which a parent filed a return in each year from 2001 to 2005, the extract adds the "
            "parents' total income from the returns for each of the five years, in constant 2005 dollars. Families in "
            "which no parent filed in one of those years are not in the history. The history holds %s families." % format(
                n_history, ",")]),
        ("4. Children's earnings, 2014 to 2022", [
            "For every child, the earnings reported on the child's own returns for each year from 2014 to 2022: "
            "employment income and net self-employment income, in current dollars. A year in which the child filed "
            "no return or reported no earnings is recorded as zero. Every child in the extract turned 34 by 2022."]),
        ("5. Dollars", [
            "Census income is in 2000 dollars, the income history in constant 2005 dollars and children's earnings in "
            "current dollars. The price index table (2005 = 100) converts any of them to constant 2005 dollars."]),
        ("6. Confidentiality", [
            "The extract is released to approved projects under the centre's agreement with the province. Results "
            "published from it must not show a count below 10, and every published figure is checked by the centre "
            "before release."]),
    ]
    for head, paras in sec:
        story.append(Paragraph(head, S["h2"]))
        for t in paras:
            story.append(Paragraph(t, S["body"]))
    doc = _doc(path, "Linked intergenerational income extract: user guide", K.CENTRE)
    doc.build(story, onFirstPage=_frame(K.CENTRE, "2026 release", "Serif"), onLaterPages=_frame(K.CENTRE, "2026 release", "Serif"))


# ---------------------------------------------------------------------------------------------- the foundation
def write_brief(path):
    import docx
    from docx.shared import Pt
    doc = docx.Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    head = doc.add_paragraph()
    r = head.add_run("%s | Programs" % K.FOUNDATION)
    r.bold = True
    r.font.size = Pt(12)
    for k, v in (("To", "Research and Evaluation"), ("From", "%s, Director of Programs" % K.PEOPLE["director"]),
                 ("Copy", "%s, Chair, Board of Trustees" % K.PEOPLE["chair"]),
                 ("Date", K.BRIEF_DATE.strftime("%-d %B %Y")),
                 ("Subject", "%s: the 2027 region, for the board on %s" % (K.PROGRAM, K.BOARD.strftime("%-d %B")))):
        p = doc.add_paragraph()
        a = p.add_run(k + ": ")
        a.bold = True
        p.add_run(v)
        p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph()
    paras = [
        "The board meets on %s to choose the region for the first three years of %s (2027 to 2029). As agreed at the "
        "September retreat, the program goes to one region: the region where children's earnings as adults depend most "
        "on their parents' long-run income." % (K.BOARD.strftime("%-d %B"), K.PROGRAM),
        "The measure is the region's intergenerational elasticity: the slope of the logarithm of a child's earnings as "
        "an adult on the logarithm of the parents' long-run income, across every child of the region. A family's "
        "long-run income is the level its yearly income moves around from one year to the next.",
        "The children are everyone born from 1984 to 1988 who lived with a parent in the region at the 2001 census, as "
        "the %s's linked extract holds them; a child counts where the parents' census income was above zero and the "
        "child had earnings in at least one year from age 30 to 34. A child's earnings as an adult are the average of "
        "the yearly earnings at ages 30 to 34 in constant 2005 dollars, a year without earnings counting as zero." % K.CENTRE,
        "The extract, its user guide, the region recode and the price index are in the folder.",
        "The board will also want to see how %s (2019 to 2023) did in each region: the young people it started there and "
        "what it spent there. The participant export taken when the program closed, the program ledger, the office list "
        "and the 2019 guidelines are in the folder with the close-out summary." % K.OLD_PROGRAM,
        "Please send the recommendation, the workbook behind it and the chart by %s so I can circulate them with the "
        "board papers." % (K.BOARD - dt.timedelta(days=7)).strftime("%-d %B"),
    ]
    for t in paras:
        p = doc.add_paragraph(t)
        p.paragraph_format.space_after = Pt(6)
    p = doc.add_paragraph(K.PEOPLE["director"])
    p.paragraph_format.space_before = Pt(8)
    doc.core_properties.author = K.PEOPLE["director"]
    doc.save(path)


def write_guidelines(path):
    S = styles("Sans", colors.HexColor("#2f5d50"))
    story = [Paragraph("%s: program guidelines" % K.OLD_PROGRAM, S["h1"]),
             Paragraph("%s. Approved by the board of trustees, 11 December 2018; in force from 7 January 2019." %
                       K.FOUNDATION, S["meta"])]
    sec = [
        ("1. What Bridges offers", [
            "A twelve-month plan for a young person aged 15 to 19: a trained mentor, a monthly bursary while the plan "
            "is followed, and help with school, training or a first job. Plans start through the regional offices; "
            "each office serves its region."]),
        ("2. Enrolment", [
            "A young person enrols at the office of the region where they live. The office opens a case record, "
            "which follows the young person to the end of the plan."]),
        ("3. Moving region", [
            "A young person who moves to another region during the plan continues it with the office there. The first "
            "office closes its record as transferred out, and the receiving office opens a record that names the "
            "record it continues."]),
        ("4. Reporting", [
            "Each year and at the program's end the foundation reports, for each region, the young people Bridges "
            "started there and what it spent there. A young person is counted once, as a start in the region of the "
            "office where they first enrolled.",
            "Spending in a region is the cost of its office's work, with the central team's costs shared among the "
            "regions by the number of young people each started. The program's accounts close when its final "
            "bursaries and evaluation are paid."]),
        ("5. Records", [
            "Case records are kept in the foundation's case management system; the ledger is kept by cost centre, one "
            "cost centre for each office and one for the central team (the program manager and coordinator, the case "
            "management system and the evaluation)."]),
    ]
    for head, paras in sec:
        story.append(Paragraph(head, S["h2"]))
        for t in paras:
            story.append(Paragraph(t, S["body"]))
    doc = _doc(path, "%s: program guidelines" % K.OLD_PROGRAM, K.FOUNDATION)
    doc.build(story, onFirstPage=_frame(K.FOUNDATION, "Bridges guidelines"), onLaterPages=_frame(K.FOUNDATION, "Bridges guidelines"))


def write_closeout(path, by_year, respondents, in_work_or_study):
    S = styles("Sans", colors.HexColor("#2f5d50"))
    story = [Paragraph("%s 2019 to 2023: close-out summary" % K.OLD_PROGRAM, S["h1"]),
             Paragraph("%s, Programs. April 2024. Prepared for the board of trustees." % K.FOUNDATION, S["meta"]),
             Paragraph("What this summary covers", S["h2"]),
             Paragraph("Bridges ran from January 2019 to December 2023 through the foundation's eight regional "
                       "offices. This summary reports what young people told us six months after their plans ended, "
                       "and the program's spending by year. Plans still running on 31 December 2023 ended with the "
                       "program, each with a final bursary. The program's accounts closed on 31 March 2024, when the "
                       "final bursaries and the evaluation were paid.", S["body"]),
             Paragraph("Six months on", S["h2"]),
             Paragraph("%s young people answered the six-month follow-up; %d per cent of them were in work, "
                       "training or study when they answered." % (format(respondents, ","), in_work_or_study), S["body"]),
             Paragraph("Spending by year", S["h2"])]
    rows = [["Year", "Spending ($)"]] + [[str(y), "{:,.0f}".format(v)] for y, v in sorted(by_year.items())]
    rows.append(["2019 to 2024", "{:,.0f}".format(sum(by_year.values()))])
    story.append(_table(S, rows, [40, 40]))
    story.append(Paragraph("Spending is the program ledger as posted, every cost centre.", S["note"]))
    story.append(Paragraph("What we learned", S["h2"]))
    for t in ["Mentors mattered most: young people who met their mentor at least twice a month finished their plans "
              "far more often.",
              "The move to the new case management system in April 2021 went smoothly for staff; the offices kept "
              "every open plan going through it.",
              "Bursaries were paid late in some months of 2020, when offices were closed; the program caught up by "
              "the autumn."]:
        story.append(Paragraph(t, S["body"]))
    doc = _doc(path, "%s 2019 to 2023: close-out summary" % K.OLD_PROGRAM, K.FOUNDATION)
    doc.build(story, onFirstPage=_frame(K.FOUNDATION, "Bridges close-out"), onLaterPages=_frame(K.FOUNDATION, "Bridges close-out"))


# ---------------------------------------------------------------------------------------------- also in the folder
def write_indicators(path):
    import openpyxl
    from openpyxl.styles import Font
    r = K.rng("indicators")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Regions 2021"
    ws.append(["Regional economic indicators, 2021 Census of Population"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Compiled by the %s from the census profile tables, March 2023." % K.FOUNDATION])
    ws.append([])
    ws.append(["Region", "Population", "Median family income 2020 ($)", "Low-income rate (%)", "Unemployment rate (%)",
               "Post-secondary credential, 25 to 64 (%)", "Median age"])
    for c in ws[4]:
        c.font = Font(bold=True)
    for name, mu in zip(K.NAMES, [row[3] for row in K.REGIONS]):
        ws.append([name, int(r.integers(90000, 1400000) if name == "Torvale Metro" else r.integers(55000, 260000)),
                   int(round(float(1.32 * 2.718281828 ** mu) / 100) * 100), round(float(r.uniform(8.5, 19.0)), 1),
                   round(float(r.uniform(5.2, 11.8)), 1), round(float(r.uniform(52, 71)), 1), round(float(r.uniform(38, 47)), 1)])
    ws.append([])
    ws.append(["Low-income rate: share of people in households below the low-income measure after tax."])
    for col, wdt in zip("ABCDEFG", (20, 12, 26, 18, 20, 34, 12)):
        ws.column_dimensions[col].width = wdt
    wb.save(path)


def write_survey(path):
    S = styles("Sans", colors.HexColor("#5a3d6b"))
    r = K.rng("aspirations")
    rows = [["Region", "Respondents", "University (%)", "College or apprenticeship (%)", "Work straight away (%)"]]
    total = 0
    for name in K.NAMES:
        u = r.uniform(34, 52)
        c = r.uniform(26, 40)
        n = int(r.integers(140, 380))
        total += n
        rows.append([name, format(n, ","), "%.0f" % u, "%.0f" % c, "%.0f" % (100 - u - c)])
    story = [Paragraph("Youth aspirations survey 2023: summary", S["h1"]),
             Paragraph("%s, Research and Evaluation. September 2023. %s young people aged 15 to 19 answered online "
                       "between April and June 2023." % (K.FOUNDATION, format(total, ",")), S["meta"]),
             Paragraph("What young people want after school", S["h2"])]
    story.append(_table(S, rows, [34, 24, 26, 40, 34]))
    story.append(Paragraph("The survey asked about plans and worries, not about family income.", S["note"]))
    story.append(Paragraph("What worries them", S["h2"]))
    story.append(Paragraph("Cost came first in every region: two in three said the cost of study or training was their "
                           "main worry, ahead of grades and of leaving home.", S["body"]))
    doc = _doc(path, "Youth aspirations survey 2023", K.FOUNDATION)
    doc.build(story, onFirstPage=_frame(K.FOUNDATION, "Aspirations survey 2023"), onLaterPages=_frame(K.FOUNDATION, "Aspirations survey 2023"))
