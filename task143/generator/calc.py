"""The golden pipeline and every rung, read from the shipped files only (the pack under target/).

    compute(target) -> dict with the features, the release model, the reserve on every route, the asks and their stops

The reserve: each release-test cell's two features from its first 100 QP-3 cycles (both cycler fleets, a continued
test stitched across its channels), the release model of record's predicted log life, the model's error on the finished
lots split into its formation-tray and within-tray parts, each Model K pack's composition (genealogy, the supplier's
tray map), and each pack's probability that its weakest cell falls below 80 per cent within 600 cycles."""
from __future__ import annotations

import csv
import datetime as dt
import itertools
import math
import os
import re
from collections import Counter, defaultdict

import numpy as np
import openpyxl
from scipy import stats

LN_W = math.log(600.0)
N_PACKS_PROMPT = 7200
ZZ = np.linspace(-9.0, 9.0, 3601)
WZ = stats.norm.pdf(ZZ)


# ---------------------------------------------------------------------------------------------------- reading
def _csv(path, delimiter=","):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh, delimiter=delimiter))


def norm_cell(x):
    """Every system's form of a Zelvora cell serial to the supplier's: ZC2638-3111526."""
    x = x.strip()
    m = re.fullmatch(r"(?:ZC)?(\d{4})[.\-]?(\d{7})", x)
    if not m:
        raise ValueError(x)
    return "ZC%s-%s" % (m.group(1), m.group(2))


def memo_model(target):
    """RL-3's coefficients and validation error as QM-07 files them."""
    from pypdf import PdfReader
    p = os.path.join(target, "qm07_release_model_rl3_zp4820.pdf")
    text = " ".join(pg.extract_text() for pg in PdfReader(p).pages)
    text = re.sub(r"\s+", " ", text)
    b0 = float(re.search(r"b0\s*=\s*([0-9]+\.[0-9]+)", text).group(1))
    b1 = -float(re.search(r"b1\s*=\s*[-−]\s*([0-9]+\.[0-9]+)", text).group(1))
    b2 = -float(re.search(r"b2\s*=\s*[-−]\s*([0-9]+\.[0-9]+)", text).group(1))
    rmse = float(re.search(r"root-mean-square error of ([0-9]+\.[0-9]+)", text).group(1))
    return {"b0": b0, "b1": b1, "b2": b2, "rmse": rmse}


def cycles(target):
    """Each cell's protocol cycles 1 to 100 from both fleets: fleet A numbers a test's first cycle 0 and writes decimal
    commas and Bielefeld time, fleet B numbers from 1 and writes UTC; a test continued on another channel is a second
    test there, so a cell's cycles are its tests' cycles in time order. Exact repeats (fleet B's overlapping pulls)
    count once."""
    rows = defaultdict(dict)
    a = _csv(os.path.join(target, "qp3_first100_cycles_fleet_a.csv"), delimiter=";")
    for r in a:
        cell = norm_cell(r["Zelle"])
        end_loc = dt.datetime.strptime(r["Zyklusende"], "%d.%m.%Y %H:%M:%S")
        rows[cell][(r["Test"], int(r["Zyklus"]))] = (end_loc, float(r["Entladekapazitaet_mAh"].replace(",", ".")),
                                                     float(r["DCIR_mOhm"].replace(",", ".")), "A")
    b = _csv(os.path.join(target, "qp3_first100_cycles_fleet_b.csv"))
    for r in b:
        cell = norm_cell(r["cell_id"])
        end_utc = dt.datetime.strptime(r["cycle_end_utc"], "%Y-%m-%dT%H:%M:%SZ")
        rows[cell][("B" + r["test_id"], int(r["cycle"]) - 1)] = (end_utc, float(r["discharge_mAh"]),
                                                                 float(r["dcir_10s_mohm"]), "B")
    out = {}
    for cell, d in rows.items():
        tests = defaultdict(list)
        for (tid, n), v in d.items():
            tests[tid].append((n, v))
        # order the tests by their first cycle's end, fleet A's local stamps put on UTC (+1 or +2 h make no difference
        # to the order: continued tests start at least a day after the first stops)
        seq = sorted(tests.values(), key=lambda lst: min(v[0] for _, v in lst))
        q, res = [], []
        for lst in seq:
            for n, v in sorted(lst):
                q.append(v[1])
                res.append(v[2])
        out[cell] = (np.array(q), np.array(res))
    return out


def features(cyc):
    f = {}
    for cell, (q, res) in cyc.items():
        assert len(q) == 100, (cell, len(q))
        f[cell] = (math.log(q[9] - q[99]), res[99] - res[9])
    return f


def register(target):
    reg = {}
    for r in _csv(os.path.join(target, "qp3_release_register.csv")):
        reg.setdefault(r["cell_serial"], r)
    return reg


def life(target):
    return {r["cell_serial"]: r for r in _csv(os.path.join(target, "qp3_life_monitoring_results.csv"))}


# ---------------------------------------------------------------------------------------------------- the model and its error
def fit_ols(feat, lifes, lots=None):
    X, y, keys = [], [], []
    for cell, L in lifes.items():
        if L["status"] != "EOL" or (lots is not None and L["lot"] not in lots):
            continue
        f1, f2 = feat[cell]
        X.append([1.0, f1, f2])
        y.append(math.log(float(L["cycles"])))
        keys.append(cell)
    X, y = np.array(X), np.array(y)
    beta, *_ = np.linalg.lstsq(X, y, rcond=None)
    return beta, X, y, keys


def lolo_rmse(feat, lifes):
    lots = sorted({L["lot"] for L in lifes.values()})
    sse, n = 0.0, 0
    for lot in lots:
        others = [l for l in lots if l != lot]
        beta, *_ = fit_ols(feat, lifes, others)
        for cell, L in lifes.items():
            if L["lot"] == lot and L["status"] == "EOL":
                f1, f2 = feat[cell]
                e = math.log(float(L["cycles"])) - (beta[0] + beta[1] * f1 + beta[2] * f2)
                sse += e * e
                n += 1
    return math.sqrt(sse / n)


def predict(m, f):
    return m["b0"] + m["b1"] * f[0] + m["b2"] * f[1]


def residuals(m, feat, lifes, reg):
    out = []
    for cell, L in lifes.items():
        if L["status"] != "EOL":
            continue
        e = math.log(float(L["cycles"])) - predict(m, feat[cell])
        out.append((L["lot"], reg[cell]["tray"], cell, e))
    return out


def components(res, by_lot=True):
    """One-way analysis of variance of the residuals on formation trays (within lots): the between-tray and the
    within-tray variance. Unbalanced trays (a test stopped before 80 per cent) take the usual n0."""
    g = defaultdict(list)
    for lot, tray, cell, e in res:
        g[(lot, tray)].append(e)
    ssw = sum(((np.array(v) - np.mean(v)) ** 2).sum() for v in g.values())
    dfw = sum(len(v) - 1 for v in g.values())
    msw = ssw / dfw
    if by_lot:
        lots = defaultdict(list)
        for (lot, tray), v in g.items():
            lots[lot].append(v)
        ssb, dfb, n_all, n2 = 0.0, 0, 0, 0.0
        for lot, vs in lots.items():
            mlot = np.mean(np.concatenate(vs))
            ssb += sum(len(v) * (np.mean(v) - mlot) ** 2 for v in vs)
            dfb += len(vs) - 1
            n_all += sum(len(v) for v in vs)
        sizes = [len(v) for v in g.values()]
        n0 = (sum(sizes) - sum(s * s for s in sizes) / sum(sizes) * 0 - 0) / len(sizes)
        n0 = (sum(sizes) - sum(sum(len(v) ** 2 for v in vs) / sum(len(v) for v in vs) for vs in lots.values())) / dfb
    else:
        allv = np.concatenate(list(g.values()))
        mall = allv.mean()
        ssb = sum(len(v) * (np.mean(v) - mall) ** 2 for v in g.values())
        dfb = len(g) - 1
        sizes = [len(v) for v in g.values()]
        n0 = (sum(sizes) - sum(s * s for s in sizes) / sum(sizes)) / dfb
    msb = ssb / dfb
    return max((msb - msw) / n0, 0.0), msw


def reml_nested(res):
    """The same split by restricted maximum likelihood with a lot level above the trays (statsmodels)."""
    import warnings
    import pandas as pd
    import statsmodels.api as sm
    df = pd.DataFrame({"lot": [r[0] for r in res], "tray": [r[0] + "|" + r[1] for r in res], "e": [r[3] for r in res]})
    md = sm.MixedLM.from_formula("e ~ 1", df, groups="lot", re_formula="1", vc_formula={"tray": "0 + C(tray)"})
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        f = md.fit(reml=True)
    return max(float(f.cov_re.iloc[0, 0]), 0.0), float(f.vcomp[0]), float(f.scale)


# ---------------------------------------------------------------------------------------------------- the packs
def genealogy(target):
    packs = defaultdict(list)
    for r in _csv(os.path.join(target, "model_k_genealogy_zc2638.csv")):
        packs[r["pack_sn"]].append((int(r["position"]), norm_cell(r["cell_scan"])))
    return {p: [c for _, c in sorted(v)] for p, v in packs.items()}


def tray_map(target):
    return {r["cell_sn"]: r["formation_tray"] for r in _csv(os.path.join(target, "zelvora_cell_data_zc2638.csv"))}


def composition(packs, tmap):
    return Counter(tuple(sorted(Counter(tmap[c] for c in cells).values(), reverse=True)) for cells in packs.values())


def q_of_z(yh, s_t, s_c, zz=ZZ, lnw=LN_W):
    """The chance that a new cell of a tray whose shared error is s_t z falls below the line, averaged over the release
    test's predicted lives."""
    yh = np.asarray(yh)
    return np.array([np.mean(stats.norm.cdf((lnw - yh - s_t * z) / s_c)) for z in zz])


def pack_probs(yh, s2t, s2c, lnw=LN_W):
    """A pack's chance by composition: three cells of one tray, two and one, one from each of three trays."""
    q = q_of_z(yh, math.sqrt(s2t), math.sqrt(s2c), lnw=lnw)
    a3 = np.trapezoid((1 - q) ** 3 * WZ, ZZ)
    a2 = np.trapezoid((1 - q) ** 2 * WZ, ZZ)
    a1 = np.trapezoid((1 - q) * WZ, ZZ)
    return {(3,): 1 - a3, (2, 1): 1 - a2 * a1, (1, 1, 1): 1 - a1 ** 3}


def reserve(comp, probs):
    return sum(n * probs[k] for k, n in comp.items())


def reserve_from(m, feat, lifes, reg, new_cells, comp, lnw=LN_W, by_lot=True):
    """The golden construction on given features and life results: (reserve, between-tray and within-tray variance)."""
    yh = np.array([predict(m, feat[c]) for c in new_cells])
    res = residuals(m, feat, lifes, reg)
    s2t, s2c = components(res, by_lot=by_lot)
    return reserve(comp, pack_probs(yh, s2t, s2c, lnw)), s2t, s2c


def independent_from(m, feat, new_cells, n_packs, sd):
    yh = np.array([predict(m, feat[c]) for c in new_cells])
    p = float(np.mean(stats.norm.cdf((LN_W - yh) / sd)))
    return n_packs * (1 - (1 - p) ** 3)


# ---------------------------------------------------------------------------------------------------- the reserve on every route
def main_call(target):
    m = memo_model(target)
    cyc = cycles(target)
    feat = features(cyc)
    reg = register(target)
    lifes = life(target)
    new_lot = "ZC2638"
    new_cells = sorted(c for c, r in reg.items() if r["lot"] == new_lot)
    yh = np.array([predict(m, feat[c]) for c in new_cells])
    res = residuals(m, feat, lifes, reg)
    s2t, s2c = components(res, by_lot=True)
    s2t_pool, s2c_pool = components(res, by_lot=False)
    packs = genealogy(target)
    tmap = tray_map(target)
    comp = composition(packs, tmap)
    n_packs = len(packs)
    probs = pack_probs(yh, s2t, s2c)
    golden = reserve(comp, probs)
    out = {"model": m, "n_new_cells": len(new_cells), "yh": yh, "s2t": s2t, "s2c": s2c, "comp": dict(comp),
           "n_packs": n_packs, "probs": probs, "golden": golden}
    e = np.array([r[3] for r in res])
    out["resid_sd_in_sample"] = float(math.sqrt((e ** 2).sum() / (len(e) - 3)))
    out["n_resid"] = len(e)
    out["icc"] = s2t / (s2t + s2c)
    # the cell chance on the default construction
    p_iid = float(np.mean(stats.norm.cdf((LN_W - yh) / m["rmse"])))
    out["p_cell"] = p_iid
    rungs = {}
    rungs["R-1 plug-in (cells predicted under 600)"] = n_packs * (1 - (1 - float(np.mean(yh < LN_W))) ** 3)
    mu, sd = float(yh.mean()), float(yh.std(ddof=1))
    rungs["R-1b plug-in normal on the predictions"] = n_packs * (1 - (1 - stats.norm.cdf((LN_W - mu) / sd)) ** 3)
    rungs["R0 default: the cell chance at the model error, three independent cells"] = n_packs * (1 - (1 - p_iid) ** 3)
    p_ins = float(np.mean(stats.norm.cdf((LN_W - yh) / out["resid_sd_in_sample"])))
    rungs["R0b the same at the in-sample residual error"] = n_packs * (1 - (1 - p_ins) ** 3)
    p_emp = float(np.mean([np.mean(v + e < LN_W) for v in yh]))
    rungs["R0c each cell with the residuals bootstrapped, independent cells"] = n_packs * (1 - (1 - p_emp) ** 3)
    # R1: the release test's own predictions checked for a tray effect, none found, so independent cells
    gn = defaultdict(list)
    for c, v in zip(new_cells, yh):
        gn[reg[c]["tray"]].append(v)
    msw_n = sum(((np.array(v) - np.mean(v)) ** 2).sum() for v in gn.values()) / sum(len(v) - 1 for v in gn.values())
    tmn = np.array([np.mean(v) for v in gn.values()])
    msb_n = 6 * ((tmn - tmn.mean()) ** 2).sum() / (len(tmn) - 1)
    out["new_pred_tray_var"] = max((msb_n - msw_n) / 6, 0.0)
    out["new_pred_icc"] = out["new_pred_tray_var"] / (out["new_pred_tray_var"] + msw_n)
    s_sh = math.sqrt(out["new_pred_tray_var"])
    q1 = np.array([np.mean(stats.norm.cdf((LN_W - yh - s_sh * z) / m["rmse"])) for z in ZZ])
    rungs["R1 the release test's own tray structure, model error independent"] = n_packs * (
        1 - np.trapezoid((1 - q1) ** 3 * WZ, ZZ))
    # R2: the finished lots' own packs, three tray-mates at a time, observed lives
    obs = defaultdict(list)
    for cell, L in lifes.items():
        if L["status"] == "EOL":
            obs[(L["lot"], reg[cell]["tray"])].append(math.log(float(L["cycles"])))
    trip = [min(v[i] for i in t) < LN_W for v in obs.values() for t in itertools.combinations(range(len(v)), 3)]
    rungs["R2 the finished lots' tray-mate packs, as observed"] = n_packs * float(np.mean(trip))
    p_hist = float(np.mean([math.log(float(L["cycles"])) < LN_W for L in lifes.values() if L["status"] == "EOL"]))
    rungs["R2b the finished lots' cell share, three independent cells"] = n_packs * (1 - (1 - p_hist) ** 3)
    lot_pred = defaultdict(list)
    for cell, L in lifes.items():
        lot_pred[L["lot"]].append(predict(m, feat[cell]))
    weak = sorted(lot_pred, key=lambda l: np.mean(lot_pred[l]))[:3]
    trip_w = [min(v[i] for i in t) < LN_W for (lot, tr), v in obs.items() if lot in weak
              for t in itertools.combinations(range(len(v)), 3)]
    rungs["R2c the three weakest finished lots' tray-mate packs"] = n_packs * float(np.mean(trip_w))
    last = sorted(lot_pred)[-1]
    p_last = float(np.mean([math.log(float(L["cycles"])) < LN_W for L in lifes.values()
                            if L["status"] == "EOL" and L["lot"] == last]))
    rungs["R2d the latest finished lot's cell share, three independent cells"] = n_packs * (1 - (1 - p_last) ** 3)
    rungs["Rover the whole pack as one cell"] = n_packs * p_iid
    rungs["Rx the cell chance compounded over two cells"] = n_packs * (1 - (1 - p_iid) ** 2)
    rungs["R3 golden: each pack's tray-shared joint chance"] = golden
    out["rungs"] = rungs
    out["weak_lots"] = weak
    # constructions of the golden method
    var = {}
    var["within-lot ANOVA, empirical predictions (golden)"] = golden
    var["trays pooled across lots"] = reserve(comp, pack_probs(yh, s2t_pool, s2c_pool))
    s2u_r, s2t_r, s2c_r = reml_nested(res)
    out["reml"] = (s2u_r, s2t_r, s2c_r)
    var["REML, lots and trays"] = reserve(comp, pack_probs(yh, s2t_r, s2c_r))
    var["REML, the lot level shared too"] = reserve(comp, pack_probs(yh, s2t_r + s2u_r, s2c_r))
    yhn = stats.norm.ppf((np.arange(400) + 0.5) / 400) * sd + mu
    var["normal fit to the predictions"] = reserve(comp, pack_probs(yhn, s2t, s2c))
    var["every pack three tray-mates"] = n_packs * probs[(3,)]
    # nonparametric: the finished lots' residuals of tray-mates three at a time, the release test's predictions
    ys = np.sort(yh)
    g = defaultdict(list)
    for lot, tray, cell, e_ in res:
        g[(lot, tray)].append(e_)
    fl = []
    for v in g.values():
        v = np.array(v)
        F = np.searchsorted(ys, LN_W - v, side="left") / len(ys)
        for t in itertools.combinations(range(len(v)), 3):
            fl.append(1 - np.prod(1 - F[list(t)]))
    var["the finished lots' tray-mate residuals resampled"] = n_packs * float(np.mean(fl))
    # Monte Carlo, a million packs
    r = np.random.default_rng(20261016)
    nmc = 1_000_000
    T = r.normal(0.0, math.sqrt(s2t), nmc)
    idx = r.integers(0, len(yh), (nmc, 3))
    C = r.normal(0.0, math.sqrt(s2c), (nmc, 3))
    var["Monte Carlo, a million single-tray packs"] = n_packs * float(np.mean((yh[idx] + T[:, None] + C).min(1) < LN_W))
    out["variants"] = var
    return out


# ---------------------------------------------------------------------------------------------------- the asks
PARTNERS = ["DVG", "TMR", "LJQ", "FRC"]
PARTNER_REGION = {"DVG": "FRB", "TMR": "UKI", "LJQ": "NOR", "FRC": "SOU"}
INV_FILE = {"DVG": "invoices_dauvergne_frb_2025-09_2026-09.xlsx", "TMR": "invoices_tolley_marsh_uki_2025-09_2026-09.xlsx",
            "LJQ": "invoices_ljungqvist_nor_2025-09_2026-09.xlsx", "FRC": "invoices_ferrucci_sou_2025-09_2026-09.xlsx"}
# each partner's labels for the invoice number, the invoice it replaces, the service period and the currency, and the
# first column of its line table
INV_LABELS = {
    "DVG": ("Facture n°", "Annule et remplace", "Période", "Devise", "N° intervention"),
    "TMR": ("Invoice", "Replaces invoice", "Service period", "Currency", "Job"),
    "LJQ": ("Faktura", "Ersätter faktura", "Period", "Valuta", "Arbetsorder"),
    "FRC": ("Fattura n.", "Sostituisce la fattura", "Periodo", "Valuta", "Ordine"),
}
MONTH_WORDS = {
    "DVG": ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre", "novembre",
            "décembre"],
    "TMR": ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"],
    "LJQ": ["januari", "februari", "mars", "april", "maj", "juni", "juli", "augusti", "september", "oktober", "november",
            "december"],
    "FRC": ["gennaio", "febbraio", "marzo", "aprile", "maggio", "giugno", "luglio", "agosto", "settembre", "ottobre",
            "novembre", "dicembre"],
}
WINDOW = [(2025, m) for m in range(10, 13)] + [(2026, m) for m in range(1, 10)]


def _xlsx_rows(path, sheet=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _d(x):
    if isinstance(x, dt.datetime):
        return x.date()
    if isinstance(x, dt.date):
        return x
    if not x:
        return None
    x = str(x)
    for f in ("%d.%m.%Y", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(x, f).date()
        except ValueError:
            pass
    raise ValueError(x)


def invoice_lines(target, keep_replaced=False, original_only=False):
    """Every partner invoice line: {(partner, job, sheet)}: case, period (year, month), labour + handling, currency, fault,
    basis. The reissued invoice stands in for the invoice it replaces."""
    out = []
    for p in PARTNERS:
        inv_l, rep_l, per_l, cur_l, job_l = INV_LABELS[p]
        wb = openpyxl.load_workbook(os.path.join(target, INV_FILE[p]), data_only=True)
        sheets, replaced = [], set()
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            meta = {r[0]: r[1] for r in rows if r and r[0] in (inv_l, rep_l, per_l, cur_l)}
            sheets.append((ws.title, meta, rows))
            if rep_l in meta:
                replaced.add(meta[rep_l])
        for title, meta, rows in sheets:
            if meta[inv_l] in replaced and not keep_replaced:
                continue
            if original_only and rep_l in meta:
                continue
            word, year = str(meta[per_l]).rsplit(" ", 1)
            month = MONTH_WORDS[p].index(word.lower()) + 1
            hi = next(i for i, r in enumerate(rows) if r and r[0] == job_l)
            for r in rows[hi + 1:]:
                if not r or r[0] is None:
                    break
                out.append(dict(partner=p, job=r[0], case=r[1], system=r[2], fault=r[3], basis=r[4],
                                amount=float(r[5]) + float(r[6]), currency=meta[cur_l], period=(int(year), month),
                                invoice=meta[inv_l], sheet=title))
    return out


def cases(target):
    return _csv(os.path.join(target, "service_cases_2024-09_2026-09.csv"))


def sap_orders(target):
    sap = _xlsx_rows(os.path.join(target, "depot_service_orders_2024-09_2026-03.xlsx"))
    ix = {h: i for i, h in enumerate(sap[0])}
    return [{h: r[i] for h, i in ix.items()} for r in sap[1:]]


def fx_table(target):
    fx = {}
    for r in _csv(os.path.join(target, "treasury_month_end_rates_2025-2026.csv")):
        d = _d(r["month_end"])
        fx[(d.year, d.month, r["currency"])] = float(r["units_per_eur"])
    return fx


def cost_tables(target):
    wb = openpyxl.load_workbook(os.path.join(target, "service_cost_sheet_fy2026.xlsx"), data_only=True)
    pack, depot, coll = [], [], None
    for row in wb["Service pack"].iter_rows(min_row=5, values_only=True):
        if row[0] == "4100-8820":
            pack.append((_d(row[2]), _d(row[3]), float(row[4])))
    for row in wb["DACH depot"].iter_rows(min_row=4, values_only=True):
        if row[0] and row[0].startswith("Battery replacement"):
            depot.append((_d(row[1]), _d(row[2]), float(row[3])))
        if row[0] and row[0].startswith("Courier collection"):
            coll = float(row[3])
    return pack, depot, coll


def valid(tab, d):
    for s_, e_, v in tab:
        if s_ <= d and (e_ is None or d <= e_):
            return v
    raise ValueError(d)


def freight(target):
    """Each returned battery's consignment, the carrier's charge for the consignment and its pieces; and the returns log's
    own freight column on the battery's row."""
    carrier = {r["consignment"]: (float(r["charge_eur"]), int(r["pieces"])) for r in
               _csv(os.path.join(target, "carrier_statement_inbound_returns_2025-09_2026-10.csv"))}
    ret = _xlsx_rows(os.path.join(target, "depot_returns_log_2024-09_2026-09.xlsx"))
    rix = {h: i for i, h in enumerate(ret[0])}
    by_ref = {}
    for r in ret[1:]:
        by_ref[str(r[rix["Reference"]])] = (str(r[rix["Consignment"]]), r[rix["Freight EUR"]])
    return carrier, by_ref


def partner_replacements(target, dating="period", inv=None, cs=None):
    """The partners' capacity replacements under the warranty, each with its repair month: the invoice's service period
    (the golden), or the case's closing month."""
    cs = cs if cs is not None else cases(target)
    inv = inv if inv is not None else invoice_lines(target)
    line = {}
    for l in inv:
        line.setdefault(l["case"], l)
    out = []
    for r in cs:
        if r["region"] == "DACH" or r["fault_code"] != "CAP" or r["warranty"] != "WTY" or r["action"] != "REPLACED":
            continue
        l = line.get(r["case_no"])
        if dating == "period":
            if l is None:
                continue
            ym = l["period"]
        else:
            if not r["closed_utc"]:
                continue
            d = _d(r["closed_utc"])
            ym = (d.year, d.month)
        out.append(dict(case=r["case_no"], partner=r["partner"], region=r["region"], ym=ym, line=l))
    return out


def ask_a(target, dating="period", freight_rule="share", fx_rule="period", inv=None, gw=False, pack_rule="dated"):
    """Each partner's capacity replacements under the warranty carried out October 2025 to September 2026: the number,
    and their average cost in euros (the invoice line in euros at the month-end rate for its service period, the
    battery's share of its return consignment's carrier charge, the service pack at its standard cost)."""
    fx = fx_table(target)
    pack, depot, coll = cost_tables(target)
    carrier, by_ref = freight(target)
    cs = cases(target)
    if gw:
        cs = [dict(r, warranty="WTY") if r["warranty"] == "GW" else r for r in cs]
    reps = partner_replacements(target, dating, inv, cs)
    inv_all = inv if inv is not None else invoice_lines(target)
    issue = {}
    res = {}
    n, tot = Counter(), defaultdict(float)
    for e in reps:
        if e["ym"] not in WINDOW:
            continue
        l = e["line"]
        if l is None:
            continue
        p = l["partner"]
        y, m = e["ym"] if fx_rule == "period" else (l["period"][0] + (l["period"][1] // 12), l["period"][1] % 12 + 1)
        if l["currency"] == "EUR":
            rate = 1.0
        else:
            while (y, m, l["currency"]) not in fx:
                y, m = (y, m - 1) if m > 1 else (y - 1, 12)
            rate = fx[(y, m, l["currency"])]
        fr = 0.0
        if e["case"] in by_ref:
            cn, row_fr = by_ref[e["case"]]
            if freight_rule == "share":
                ch, pcs = carrier[cn]
                fr = ch / pcs
            elif freight_rule == "row":
                fr = float(row_fr or 0.0)
        d = dt.date(e["ym"][0], e["ym"][1], 15)
        pc = valid(pack, d) if pack_rule == "dated" else pack[-1][2]
        c = l["amount"] / rate + fr + pc
        n[p] += 1
        tot[p] += c
    return {p: (n[p], tot[p] / n[p] if n[p] else 0.0) for p in PARTNERS}


def depot_replacements(target, dedupe=True):
    """The DACH depot's capacity replacements under the warranty and their repair days: SAP orders completed in SAP
    (Techn. Abschluss), and the depot's cases in the service case system, those carried over from SAP and completed
    there counted once."""
    sap = sap_orders(target)
    out, done = [], set()
    for r in sap:
        if r["Schadenscode"] == "B10" and r["Auftragsart"] == "ZS01" and r["Techn. Abschluss"] and r["Einbau-SN"]:
            d = _d(r["Techn. Abschluss"])
            out.append(("SAP", str(r["Auftrag"]), d))
            done.add(str(r["Auftrag"]))
    for r in cases(target):
        if r["region"] != "DACH" or r["fault_code"] != "CAP" or r["warranty"] != "WTY" or r["action"] != "REPLACED":
            continue
        if not r["closed_utc"]:
            continue
        if dedupe and r["legacy_ref"] and r["legacy_ref"] in done:
            continue
        out.append(("CASE", r["case_no"], _d(r["closed_utc"])))
    return out


def ask_b(target, dating="period", dedupe=True, gw=False):
    """The capacity replacements under the warranty carried out in each month, October 2025 to September 2026, all five
    regions: the depot's by their repair day, the partners' by their invoice's service period."""
    cnt = Counter()
    for src, ref, d in depot_replacements(target, dedupe):
        cnt[(d.year, d.month)] += 1
    cs = cases(target)
    if gw:
        cs = [dict(r, warranty="WTY") if r["warranty"] == "GW" else r for r in cs]
    for e in partner_replacements(target, dating, cs=cs):
        cnt[e["ym"]] += 1
    return {ym: cnt[ym] for ym in WINDOW}


def ask_b_by_region(target):
    out = defaultdict(Counter)
    for src, ref, d in depot_replacements(target):
        out[(d.year, d.month)]["DACH"] += 1
    for e in partner_replacements(target):
        out[e["ym"]][e["region"]] += 1
    return {ym: dict(out[ym]) for ym in WINDOW}


def ask_stops(target):
    """Every complete pipeline that misses: each mishandles one thing and handles the rest as the golden does."""
    inv_both = invoice_lines(target, keep_replaced=True)
    inv_orig = invoice_lines(target, keep_replaced=True, original_only=True)
    A = {}
    A["golden"] = ask_a(target)
    A["A1 the returns log's freight on each battery's row"] = ask_a(target, freight_rule="row")
    A["A2 each case dated by its closing day"] = ask_a(target, dating="closed")
    A["A3 Tolley Marsh's March invoice and its reissue both"] = ask_a(target, inv=inv_both)
    A["A4 Tolley Marsh's March invoice as first sent"] = ask_a(target, inv=inv_orig)
    A["A5 each rate taken for the month of the invoice date"] = ask_a(target, fx_rule="issue")
    A["A6 goodwill replacements kept"] = ask_a(target, gw=True)
    A["A7 no freight"] = ask_a(target, freight_rule="none")
    A["A8 the current pack cost throughout"] = ask_a(target, pack_rule="current")
    A["cracker: closing days and the log's freight"] = ask_a(target, dating="closed", freight_rule="row")
    B = {}
    B["golden"] = ask_b(target)
    B["B1 the partners' cases by their closing month"] = ask_b(target, dating="closed")
    B["B2 the carried-over depot orders counted in both systems"] = ask_b(target, dedupe=False)
    B["B3 goodwill replacements kept"] = ask_b(target, gw=True)
    B["cracker: closing months"] = ask_b(target, dating="closed")
    return A, B


def compute(target):
    c = main_call(target)
    c["ask_a"] = ask_a(target)
    c["ask_b"] = ask_b(target)
    c["ask_b_region"] = ask_b_by_region(target)
    A, B = ask_stops(target)
    c["ask_a_stops"] = A
    c["ask_b_stops"] = B
    return c
