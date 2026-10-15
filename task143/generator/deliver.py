"""The golden deliverables: z2638_warranty_reserve.xlsx (the reserve and how it is built, the four partners' capacity
replacements under the warranty and their average cost, the capacity replacements carried out in each month by service
region) and z2638_reserve_note.docx (the one-page note to Product Finance opening with the reserve, with the chart of
the monthly replacements), plus generator/figures.json. Every figure is computed from the shipped pack (calc.py)."""
from __future__ import annotations

import io
import json
import math
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import calc as CA
import common as K
import postprocess as PP

NOTE_DATE = "15 October 2026"
PARTNER_NAME = {"DVG": "Dauvergne Réparation (FRB)", "TMR": "Tolley Marsh Repair (UKI)",
                "LJQ": "Ljungqvist Service (NOR)", "FRC": "Ferrucci Assistenza (SOU)"}
REGIONS = ["DACH", "FRB", "UKI", "NOR", "SOU"]
MONTH_NAME = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
              "November", "December"]


def f0(x):
    return "{:,}".format(int(x))


def pct1(x):
    return "%.1f" % K.round_dp(100.0 * x, 1)


def month_label(ym):
    return "%s %d" % (MONTH_NAME[ym[1] - 1], ym[0])


def compute(target):
    c = CA.compute(target)
    probs = c["probs"]
    comp = c["comp"]
    c["reserve"] = K.round_to(c["golden"], 50)
    c["pack_chance"] = c["golden"] / c["n_packs"]
    c["indep"] = c["rungs"]["R0 default: the cell chance at the model error, three independent cells"]
    c["indep_chance"] = 1 - (1 - c["p_cell"]) ** 3
    c["comp_rows"] = [(k, comp.get(k, 0), probs[k]) for k in ((3,), (2, 1), (1, 1, 1))]
    # the partners' cost parts, each averaged over the partner's replacements
    c["ask_a_parts"] = ask_a_parts(target)
    return c


def ask_a_parts(target):
    """The average invoice line in euros, freight share and service pack of each partner's replacements (the golden's
    own arithmetic, split by part)."""
    fx = CA.fx_table(target)
    pack, depot, coll = CA.cost_tables(target)
    carrier, by_ref = CA.freight(target)
    out = {}
    sums = {p: [0.0, 0.0, 0.0, 0] for p in CA.PARTNERS}
    for e in CA.partner_replacements(target):
        if e["ym"] not in CA.WINDOW or e["line"] is None:
            continue
        l = e["line"]
        y, m = e["ym"]
        rate = 1.0 if l["currency"] == "EUR" else fx[(y, m, l["currency"])]
        fr = 0.0
        if e["case"] in by_ref:
            ch, pcs = carrier[by_ref[e["case"]][0]]
            fr = ch / pcs
        s = sums[l["partner"]]
        s[0] += l["amount"] / rate
        s[1] += fr
        s[2] += CA.valid(pack, CA.dt.date(y, m, 15))
        s[3] += 1
    for p, (a, b, d, n) in sums.items():
        out[p] = (a / n, b / n, d / n, n)
    return out


def figures(c):
    m = c["model"]
    out = {
        "call": {"reserve_packs": f0(c["reserve"]), "lot": "ZC2638", "packs": f0(c["n_packs"])},
        "components": {
            "expected_packs_unrounded": "%.1f" % K.round_dp(c["golden"], 1),
            "pack_chance_pct": pct1(c["pack_chance"]),
            "packs_three_traymates": f0(c["comp"].get((3,), 0)),
            "packs_two_and_one": f0(c["comp"].get((2, 1), 0)),
            "packs_three_trays": f0(c["comp"].get((1, 1, 1), 0)),
            "tray_sd": "%.3f" % K.round_dp(math.sqrt(c["s2t"]), 3),
            "cell_sd": "%.3f" % K.round_dp(math.sqrt(c["s2c"]), 3),
            "tray_share_pct": "%d" % K.round_to(100 * c["icc"], 1),
        },
        "steps": {
            "release_cells": f0(c["n_new_cells"]),
            "release_trays": "20",
            "finished_lots": "20",
            "life_test_cells_eol": f0(c["n_resid"]),
            "rl3_b0": "%.4f" % m["b0"], "rl3_b1": "%.4f" % m["b1"], "rl3_b2": "%.4f" % m["b2"],
            "rl3_rmse": "%.3f" % m["rmse"],
            "cell_chance_pct": pct1(c["p_cell"]),
            "release_test_tray_share_pct": "%d" % K.round_to(100 * c["new_pred_icc"], 1),
            "warranty_cycles": "600", "eol_pct": "80", "eol_mah": "3,856",
        },
        "filed": {"fin_rs04_clauses": ["2.2", "3.2", "3.3", "3.4"], "spec_section": "4", "qm07_sections": ["3", "5"]},
        "rungs": {"finished_lots_traymate_packs": f0(K.round_to(
            c["rungs"]["R2 the finished lots' tray-mate packs, as observed"], 50))},
        "default_path": {
            "independent_pack_chance_pct": pct1(c["indep_chance"]),
            "independent_reserve_unrounded": f0(K.round_to(c["indep"], 1)),
            "independent_reserve_packs": f0(K.round_to(c["indep"], 50)),
        },
        "ask_a": {p: {"replacements": f0(c["ask_a"][p][0]), "average_cost_eur": f0(K.round_to(c["ask_a"][p][1], 1))}
                  for p in CA.PARTNERS},
        "ask_a_parts": {p: {"invoice_eur": "%.2f" % K.round_dp(v[0], 2), "freight_eur": "%.2f" % K.round_dp(v[1], 2),
                            "pack_eur": "%.2f" % K.round_dp(v[2], 2)} for p, v in c["ask_a_parts"].items()},
        "ask_b": {month_label(ym): f0(c["ask_b"][ym]) for ym in CA.WINDOW},
        "ask_b_total": f0(sum(c["ask_b"].values())),
        "ask_b_region": {month_label(ym): {r: f0(c["ask_b_region"][ym].get(r, 0)) for r in REGIONS} for ym in CA.WINDOW},
    }
    return out


# ---------------------------------------------------------------------------------------------------- the workbook
def _title(ws, title, sub):
    ws["A1"] = title
    ws["A1"].font = Font(name="Calibri", bold=True, size=13, color="0F4C5C")
    ws["A2"] = sub
    ws["A2"].font = Font(name="Calibri", italic=True, size=9.5, color="595959")


def _head(ws, row, values, widths=None):
    for k, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=k, value=v)
        cell.font = Font(name="Calibri", bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE8EC")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    if widths:
        for k, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(k)].width = w


def _row(ws, row, values, fmts=None, bold=False):
    for k, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=k, value=v)
        cell.font = Font(name="Calibri", bold=bold)
        if fmts and fmts[k - 1]:
            cell.number_format = fmts[k - 1]


def _note(ws, row, text, ncol, height=30):
    ws.cell(row=row, column=1, value=text)
    ws.cell(row=row, column=1).font = Font(name="Calibri", italic=True, size=9, color="595959")
    ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
    ws.row_dimensions[row].height = height


def write_book(path, c):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reserve"
    _title(ws, "Lot ZC2638: reserve %s Model K packs at release" % f0(K.round_to(c["golden"], 50)),
           "FIN-RS-04: the expected number of the lot's Model K packs whose full charge capacity falls below 80 per "
           "cent of rated capacity within their first 600 cycles. Release Friday 16 October 2026.")
    _head(ws, 4, ["Item", "Value"], [62, 16])
    rows = [
        ("Reserve to book (packs, nearest 50)", c["reserve"], "#,##0", True),
        ("Expected packs below 80 per cent within 600 cycles (unrounded)", K.round_dp(c["golden"], 1), "#,##0.0", False),
        ("Model K packs built from ZC2638 (genealogy)", c["n_packs"], "#,##0", False),
        ("Expected share of the packs (%)", float(pct1(c["pack_chance"])), "0.0", False),
    ]
    r = 5
    for label, v, fmt, b in rows:
        _row(ws, r, [label, v], [None, fmt], bold=b)
        r += 1
    r += 1
    _head(ws, r, ["Pack composition (genealogy and Zelvora's tray map)", "Packs", "Chance a pack falls below 80% (%)",
                  "Expected packs"], [62, 16, 18, 16])
    r += 1
    labels = {(3,): "Three cells from one formation tray", (2, 1): "Two cells from one tray, one from another",
              (1, 1, 1): "Three cells from three trays"}
    for k, n, pr in c["comp_rows"]:
        _row(ws, r, [labels[k], n, K.round_dp(100 * pr, 2), K.round_dp(n * pr, 1)], [None, "#,##0", "0.00", "#,##0.0"])
        r += 1
    r += 1
    _head(ws, r, ["RL-3's error on the finished lots (log cycle life, 20 lots, %s cells)" % f0(c["n_resid"]),
                  "Standard deviation", "Variance", "Share (%)"], [62, 16, 18, 16])
    r += 1
    tot = c["s2t"] + c["s2c"]
    _row(ws, r, ["Shared by the cells of a formation tray", K.round_dp(math.sqrt(c["s2t"]), 3),
                 K.round_dp(c["s2t"], 4), K.round_to(100 * c["s2t"] / tot, 1)], [None, "0.000", "0.0000", "0"])
    r += 1
    _row(ws, r, ["A cell's own, within its tray", K.round_dp(math.sqrt(c["s2c"]), 3), K.round_dp(c["s2c"], 4),
                 K.round_to(100 * c["s2c"] / tot, 1)], [None, "0.000", "0.0000", "0"])
    r += 2
    _head(ws, r, ["For comparison", "Value"], [62, 16])
    r += 1
    _row(ws, r, ["A release-test cell's chance at RL-3's validation error of %.3f (%%)" % c["model"]["rmse"],
                 float(pct1(c["p_cell"]))], [None, "0.0"])
    r += 1
    _row(ws, r, ["The three cells of a pack taken as independent draws: a pack's chance (%)",
                 float(pct1(c["indep_chance"]))], [None, "0.0"])
    r += 1
    _row(ws, r, ["The three cells of a pack taken as independent draws: packs", K.round_to(c["indep"], 1)],
         [None, "#,##0"])
    r += 2
    _note(ws, r, "Each release-test cell's log life predicted by RL-3 (QM-07 rev. 4) from its first 100 QP-3 cycles. "
          "RL-3's error on the finished lots' life tests split by one-way analysis of variance within lots into the "
          "part a formation tray's cells share and each cell's own. A pack's chance is the expectation, over its "
          "tray's shared error, of the chance that any of its cells falls below 80 per cent within 600 cycles.", 4, 58)
    ws.freeze_panes = "A5"
    # Ask A
    ws = wb.create_sheet("Partner cost")
    _title(ws, "Capacity replacements under the warranty carried out by the repair partners, October 2025 to "
           "September 2026", "Per replacement: the partner's invoice line in euros at treasury's month-end rate for "
           "the invoice's service period, the removed battery's share of its consignment's carrier charge, and the "
           "service pack at standard cost.")
    _head(ws, 4, ["Partner (region)", "Replacements", "Invoice line (EUR)", "Return freight (EUR)",
                  "Service pack (EUR)", "Average cost (EUR, nearest euro)"], [30, 14, 14, 14, 14, 18])
    r = 5
    for p in CA.PARTNERS:
        n, avg = c["ask_a"][p]
        a, b, d, nn = c["ask_a_parts"][p]
        _row(ws, r, [PARTNER_NAME[p], n, K.round_dp(a, 2), K.round_dp(b, 2), K.round_dp(d, 2), K.round_to(avg, 1)],
             [None, "#,##0", "#,##0.00", "#,##0.00", "#,##0.00", "#,##0"])
        r += 1
    r += 1
    _note(ws, r, "Carried out in the month of the invoice's service period. Goodwill replacements after 600 cycles "
          "and diagnosis-only cases are not warranty replacements; a battery that never reached Bielefeld carries no "
          "freight. Tolley Marsh Repair's March 2026 invoice is taken as reissued (TM-2603-R1). Parts are averages to "
          "the "
          "cent; the average cost is rounded once from unrounded working.", 6, 44)
    ws.freeze_panes = "A5"
    # Ask B
    ws = wb.create_sheet("Monthly replacements")
    _title(ws, "Capacity replacements under the warranty carried out each month, October 2025 to September 2026",
           "DACH depot by its repair day (SAP Techn. Abschluss to 1 March 2026, the case system after); the partners "
           "by their invoices' service periods.")
    _head(ws, 4, ["Month"] + REGIONS + ["All regions"], [18, 9, 9, 9, 9, 9, 12])
    r = 5
    for ym in CA.WINDOW:
        reg = c["ask_b_region"][ym]
        _row(ws, r, [month_label(ym)] + [reg.get(x, 0) for x in REGIONS] + [c["ask_b"][ym]],
             [None] + ["#,##0"] * 6)
        r += 1
    _row(ws, r, ["Twelve months"] + [sum(c["ask_b_region"][ym].get(x, 0) for ym in CA.WINDOW) for x in REGIONS] +
         [sum(c["ask_b"].values())], [None] + ["#,##0"] * 6, bold=True)
    r += 2
    _note(ws, r, "DACH orders the depot began in SAP and completed in SAP are counted once, from SAP. A partner's "
          "case closes in the case system after the month its invoice bills it in, so its closing month is not the "
          "month of the repair.", 7, 32)
    ws.freeze_panes = "A5"
    ws = wb.create_sheet("Notes")
    _title(ws, "Notes", "Product Finance, Notebook Division. Prepared %s for the release of lot ZC2638." % NOTE_DATE)
    notes = [
        ("Sources", "QM-07 rev. 4; the fleet A and fleet B cycler exports of the first 100 QP-3 cycles; the release "
                    "test register; the life test results of lots ZC2437 to ZC2614; Zelvora's cell data package and "
                    "the line genealogy for ZC2638 (pulled 12 October 2026); the service case system and the DACH "
                    "depot's SAP orders (5 October 2026); the returns depot log, the carrier statement, the four "
                    "partners' invoices, treasury's month-end rates and the FY2026 service cost sheet."),
        ("Reserve", "FIN-RS-04 rev. 3, clauses 2.2, 3.2 and 3.3: every Model K pack built from the lot, a pack below "
                    "80 per cent when its weakest cell is (pack specification rev. D, section 4)."),
        ("Rounding", "The reserve to the nearest 50 packs, costs to the nearest euro, each rounded once from unrounded "
                     "working."),
    ]
    _head(ws, 4, ["Item", "Note"], [14, 112])
    for k, (a, b) in enumerate(notes):
        _row(ws, 5 + k, [a, b])
        ws.cell(row=5 + k, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[5 + k].height = 58 if len(b) > 300 else 30
    for w in wb.worksheets:
        w.page_setup.orientation = "landscape"
        w.page_setup.paperSize = w.PAPERSIZE_A4
        w.sheet_properties.pageSetUpPr.fitToPage = True
        w.page_setup.fitToWidth = 1
        w.page_setup.fitToHeight = 0
    wb.save(path)
    PP.normalize_ooxml(path, "Product Finance", K.PEOPLE["controller"], "2026-10-14T11:08:00Z", "2026-10-15T15:42:00Z",
                       excel=True, company=K.COMPANY)


# ---------------------------------------------------------------------------------------------------- the chart
COLORS = {"DACH": "#0F4C5C", "FRB": "#5F9EA0", "UKI": "#E0A458", "NOR": "#9C6644", "SOU": "#B8C4BB"}


def chart_png(c):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    matplotlib.rcParams["svg.hashsalt"] = "task143"
    matplotlib.rcParams["font.family"] = "DejaVu Sans"
    labels = ["%s %02d" % (MONTH_NAME[ym[1] - 1][:3], ym[0] % 100) for ym in CA.WINDOW]
    fig, ax = plt.subplots(figsize=(7.6, 3.6), dpi=150)
    base = [0] * len(CA.WINDOW)
    for reg in REGIONS:
        vals = [c["ask_b_region"][ym].get(reg, 0) for ym in CA.WINDOW]
        ax.bar(labels, vals, bottom=base, color=COLORS[reg], width=0.68, label=reg, edgecolor="white", linewidth=0.4)
        base = [b + v for b, v in zip(base, vals)]
    for k, ym in enumerate(CA.WINDOW):
        ax.text(k, base[k] + 1.0, str(c["ask_b"][ym]), ha="center", va="bottom", fontsize=7.5, color="#1D2229")
    ax.set_ylim(0, max(base) * 1.16)
    ax.set_ylabel("Replacements", fontsize=8.5)
    ax.tick_params(axis="both", labelsize=7.5)
    first, last = CA.WINDOW[0], CA.WINDOW[-1]
    ax.set_title("Capacity replacements under the warranty climbed through the year\n%d in %s, %d in %s, by "
                 "service region" % (c["ask_b"][first], month_label(first), c["ask_b"][last], month_label(last)),
                 fontsize=9, loc="left")
    ax.legend(ncol=5, fontsize=7.2, frameon=False, loc="upper left")
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    fig.text(0.01, 0.01, "DACH by repair day (SAP, then the case system); partners by invoice service period. "
             "October 2025 to September 2026, %s in all." % f0(sum(c["ask_b"].values())), fontsize=6.5,
             color="#5A6270")
    fig.tight_layout(rect=(0, 0.04, 1, 1))
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


# ---------------------------------------------------------------------------------------------------- the note
def _page_field(par):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        r = par.add_run()
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        elif text.strip() == "PAGE":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            r._r.append(it)
        else:
            r.text = text


def write_note(path, c):
    import docx
    from docx.shared import Cm, Pt, RGBColor
    doc = docx.Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    st.paragraph_format.space_after = Pt(4)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.5)
    fp = sec.footer.paragraphs[0]
    fp.text = "Kalvenhorst Systeme GmbH, Product Finance, Notebook Division. Page "
    _page_field(fp)
    for r in fp.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x5A, 0x62, 0x70)
    p = doc.add_paragraph()
    r = p.add_run("Lot ZC2638: reserve %s Model K packs at release" % f0(K.round_to(c["golden"], 50)))
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(0x0F, 0x4C, 0x5C)
    for a_, b_ in (("To:", "%s, Product Finance Controller, Notebook Division" % K.PEOPLE["controller"]),
                   ("Copy:", "%s, Battery Reliability Lab" % K.PEOPLE["lab"]),
                   ("From:", "Product Finance, Notebook Division"),
                   ("Date:", NOTE_DATE + ", for the release on Friday 16 October"),
                   ("Subject:", "Warranty reserve for the Model K packs of Zelvora lot ZC2638 (FIN-RS-04)")):
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(0)
        q.add_run(a_ + " ").bold = True
        q.add_run(b_)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run("Book a reserve of %s packs for the %s Model K packs built from lot ZC2638." % (
        f0(c["reserve"]), f0(c["n_packs"])))
    r.bold = True
    r.font.size = Pt(11.5)
    doc.add_paragraph(
        "The expected number below 80 per cent of rated capacity within 600 cycles is %s, %s per cent of the packs. "
        "The 120 cells of the release test predict a chance of %s per cent for a cell at RL-3's validation error. A "
        "pack falls below 80 per cent with its weakest cell, and the line builds its packs from one formation tray at "
        "a time: %s of the %s packs are three tray-mates, %s take two cells from one tray and one from another, and %s "
        "a cell from each of three trays." % (
            "%.1f" % K.round_dp(c["golden"], 1), pct1(c["pack_chance"]), pct1(c["p_cell"]),
            f0(c["comp"].get((3,), 0)), f0(c["n_packs"]), f0(c["comp"].get((2, 1), 0)),
            "one pack has" if c["comp"].get((1, 1, 1), 0) == 1 else f0(c["comp"].get((1, 1, 1), 0)) + " packs have"))
    doc.add_paragraph(
        "Most of RL-3's error is the tray's, not the cell's. On the twenty finished lots, %s cells followed to 80 per "
        "cent, the error of the log life has a standard deviation of %s shared by a tray's cells and %s of each cell's "
        "own: %s per cent of its variance sits at the tray. The first 100 cycles do not show it (the release test's own "
        "predictions carry almost no tray structure); only whole trays cycled to the end do. The three cells of a pack "
        "therefore succeed or fail largely together, so a pack's chance is the chance, over its tray's shared error, "
        "that any of its cells falls short. Compounding the cell chance over three independent cells would give %s "
        "per cent and %s packs." % (
            f0(c["n_resid"]), "%.3f" % K.round_dp(math.sqrt(c["s2t"]), 3), "%.3f" % K.round_dp(math.sqrt(c["s2c"]), 3),
            "%d" % K.round_to(100 * c["icc"], 1), pct1(c["indep_chance"]), f0(K.round_to(c["indep"], 1))))
    q = doc.add_paragraph()
    q.paragraph_format.space_before = Pt(4)
    q.add_run("Cost of a capacity replacement at the repair partners, October 2025 to September 2026").bold = True
    tb = doc.add_table(rows=1, cols=4)
    tb.style = "Table Grid"
    for k, x in enumerate(["Partner (region)", "Replacements", "Average cost (EUR)", "Of which return freight (EUR)"]):
        tb.rows[0].cells[k].text = x
    for p_ in CA.PARTNERS:
        n, avg = c["ask_a"][p_]
        row = tb.add_row().cells
        row[0].text = PARTNER_NAME[p_]
        row[1].text = f0(n)
        row[2].text = f0(K.round_to(avg, 1))
        row[3].text = "%.2f" % K.round_dp(c["ask_a_parts"][p_][1], 2)
    for rr in tb.rows:
        for cell in rr.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                for run in para.runs:
                    run.font.size = Pt(9)
    q = doc.add_paragraph()
    r = q.add_run("Invoice line at treasury's month-end rate for its service period, the battery's share of its "
                  "consignment's carrier charge, the service pack at standard cost. The DACH depot's internal cost is "
                  "in the cost sheet.")
    r.italic = True
    r.font.size = Pt(8)
    doc.add_picture(io.BytesIO(chart_png(c)), width=Cm(16.5))
    # the picture's names as Word gives them
    pic = doc.inline_shapes[0]._inline
    pic.docPr.set("name", "Picture 1")
    pic.docPr.set("descr", "Capacity replacements under the warranty each month by service region")
    pic.graphic.graphicData.pic.nvPicPr.cNvPr.set("id", "1")
    pic.graphic.graphicData.pic.nvPicPr.cNvPr.set("name", "Picture 1")
    q = doc.add_paragraph()
    r = q.add_run("%s capacity replacements under the warranty in the twelve months, %s in September 2026. "
                  "z2638_warranty_reserve.xlsx has the reserve's working, the partners' costs and the monthly table."
                  % (f0(sum(c["ask_b"].values())), f0(c["ask_b"][(2026, 9)])))
    r.font.size = Pt(9)
    q = doc.add_paragraph()
    q.paragraph_format.space_before = Pt(4)
    r = q.add_run("Packs: every Model K pack the battery plant built from the lot's cells by the release day, as "
                  "FIN-RS-04 clause 2.2 counts them, from the line genealogy. Cycle lives are on QP-3, which the "
                  "pack's gauge counts as its cycles (clause 3.3). Tolley Marsh Repair's March 2026 invoice is "
                  "taken as reissued (TM-2603-R1).")
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor(0x5A, 0x62, 0x70)
    doc.core_properties.author = "Product Finance"
    doc.core_properties.title = "Warranty reserve ZC2638"
    doc.save(path)
    PP.normalize_docx(path, "Product Finance", K.PEOPLE["controller"], "2026-10-15T09:12:00Z", "2026-10-15T16:05:00Z",
                      revision=4, total_time=96, seed="task143/note", title="Warranty reserve ZC2638",
                      company=K.COMPANY)


def main(target, golden, figures_path):
    os.makedirs(golden, exist_ok=True)
    c = compute(target)
    figs = figures(c)
    write_book(os.path.join(golden, K.GOLD_BOOK), c)
    write_note(os.path.join(golden, K.GOLD_NOTE), c)
    with open(figures_path, "w") as fh:
        json.dump(figs, fh, indent=2, sort_keys=False, ensure_ascii=False)
        fh.write("\n")
    return c, figs
