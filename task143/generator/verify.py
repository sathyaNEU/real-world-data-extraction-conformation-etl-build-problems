"""The independent verifier: reads only the shipped bundle (target/) and the golden workbook, on its own code path (pandas
readers, its own feature extraction, the release model refitted, a nested mixed model by restricted maximum likelihood
for the error's tray and cell parts, Gauss-Hermite quadrature over the tray's shared error, the asks by joins), and
recomputes the reserve, every partner's replacements and average cost, and the month counts."""
from __future__ import annotations

import math
import os
import re
import sys
import warnings

import numpy as np
import openpyxl
import pandas as pd
from scipy import stats


def cell_key(s):
    m = re.fullmatch(r"(?:ZC)?(\d{4})[.\-]?(\d{7})", str(s).strip())
    return "ZC%s-%s" % (m.group(1), m.group(2))


def features(t):
    a = pd.read_csv(os.path.join(t, "qp3_first100_cycles_fleet_a.csv"), sep=";", dtype=str)
    a = pd.DataFrame({"cell": a["Zelle"].map(cell_key), "test": "A" + a["Test"], "n": a["Zyklus"].astype(int) + 1,
                      "q": a["Entladekapazitaet_mAh"].str.replace(",", ".").astype(float),
                      "r": a["DCIR_mOhm"].str.replace(",", ".").astype(float),
                      "t": pd.to_datetime(a["Zyklusende"], format="%d.%m.%Y %H:%M:%S")})
    b = pd.read_csv(os.path.join(t, "qp3_first100_cycles_fleet_b.csv"), dtype=str)
    b = pd.DataFrame({"cell": b["cell_id"].map(cell_key), "test": "B" + b["test_id"], "n": b["cycle"].astype(int),
                      "q": b["discharge_mAh"].astype(float), "r": b["dcir_10s_mohm"].astype(float),
                      "t": pd.to_datetime(b["cycle_end_utc"], format="%Y-%m-%dT%H:%M:%SZ")})
    b = b.drop_duplicates()
    x = pd.concat([a, b], ignore_index=True)
    first = x.groupby(["cell", "test"]).t.min().rename("t0")
    x = x.join(first, on=["cell", "test"]).sort_values(["cell", "t0", "n"])
    x["k"] = x.groupby("cell").cumcount() + 1          # the protocol's cycle across a continued test
    assert x.groupby("cell").k.max().eq(100).all()
    q = x.pivot(index="cell", columns="k", values="q")
    r = x.pivot(index="cell", columns="k", values="r")
    return pd.DataFrame({"f1": np.log(q[10] - q[100]), "f2": r[100] - r[10]})


def memo(t):
    from pypdf import PdfReader
    s = re.sub(r"\s+", " ", " ".join(p.extract_text() for p in PdfReader(os.path.join(
        t, "qm07_release_model_rl3_zp4820.pdf")).pages))
    g = lambda pat: float(re.search(pat, s).group(1))
    return (g(r"b0\s*=\s*([0-9]+\.[0-9]+)"), -g(r"b1\s*=\s*[-−]\s*([0-9]+\.[0-9]+)"),
            -g(r"b2\s*=\s*[-−]\s*([0-9]+\.[0-9]+)"))


def reserve(t):
    f = features(t)
    reg = pd.read_csv(os.path.join(t, "qp3_release_register.csv"), dtype=str).drop_duplicates("cell_serial")
    reg = reg.set_index("cell_serial")
    life = pd.read_csv(os.path.join(t, "qp3_life_monitoring_results.csv"), dtype=str)
    life = life[life.status == "EOL"].set_index("cell_serial")
    b0, b1, b2 = memo(t)
    d = f.join(life[["lot", "cycles"]], how="inner").join(reg[["tray"]])
    X = np.column_stack([np.ones(len(d)), d.f1, d.f2])
    beta = np.linalg.lstsq(X, np.log(d.cycles.astype(float)), rcond=None)[0]
    assert np.allclose(beta, [b0, b1, b2], atol=6e-4), (beta, (b0, b1, b2))
    d["e"] = np.log(d.cycles.astype(float)) - (b0 + b1 * d.f1 + b2 * d.f2)
    d["lt"] = d.lot + "/" + d.tray
    import statsmodels.formula.api as smf
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        fit = smf.mixedlm("e ~ 1", d, groups="lt").fit(reml=True)
    s2t, s2c = float(fit.cov_re.iloc[0, 0]), float(fit.scale)
    new = reg[reg.lot == "ZC2638"].index
    yh = (b0 + b1 * f.loc[new, "f1"] + b2 * f.loc[new, "f2"]).to_numpy()
    gen = pd.read_csv(os.path.join(t, "model_k_genealogy_zc2638.csv"), dtype=str)
    cells = pd.read_csv(os.path.join(t, "zelvora_cell_data_zc2638.csv"), dtype=str).set_index("cell_sn")
    gen["tray"] = gen.cell_scan.map(cell_key).map(cells.formation_tray)
    assert gen.tray.notna().all()
    kinds = gen.groupby("pack_sn").tray.agg(lambda s: tuple(sorted(s.value_counts().tolist(), reverse=True)))
    z, w = np.polynomial.hermite_e.hermegauss(80)
    w = w / w.sum()
    lnw = math.log(600.0)
    q = np.array([stats.norm.cdf((lnw - yh - math.sqrt(s2t) * zz) / math.sqrt(s2c)).mean() for zz in z])
    a1, a2, a3 = (w * (1 - q)).sum(), (w * (1 - q) ** 2).sum(), (w * (1 - q) ** 3).sum()
    p = {(3,): 1 - a3, (2, 1): 1 - a2 * a1, (1, 1, 1): 1 - a1 ** 3}
    return sum(p[k] for k in kinds), len(kinds), s2t, s2c


MONTHS = {"janvier": 1, "février": 2, "mars": 3, "avril": 4, "mai": 5, "juin": 6, "juillet": 7, "août": 8,
          "septembre": 9, "octobre": 10, "novembre": 11, "décembre": 12, "jan": 1, "feb": 2, "mar": 3, "apr": 4,
          "may": 5, "jun": 6, "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12, "januari": 1,
          "februari": 2, "april": 4, "maj": 5, "juni": 6, "juli": 7, "augusti": 8, "september": 9, "oktober": 10,
          "november": 11, "december": 12, "gennaio": 1, "febbraio": 2, "marzo": 3, "aprile": 4, "maggio": 5,
          "giugno": 6, "luglio": 7, "agosto": 8, "settembre": 9, "ottobre": 10, "dicembre": 12}


def invoice_lines(t):
    out = []
    for f in sorted(os.listdir(t)):
        if not f.startswith("invoices_"):
            continue
        wb = openpyxl.load_workbook(os.path.join(t, f), data_only=True)
        sheets = []
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            head = {str(r[0]): r[1] for r in rows[:12] if r and r[0] is not None and len(r) > 1}
            number = head[[k for k in head if re.match(r"(Invoice|Facture|Faktura|Fattura)", k)][0]]
            repl = [v for k, v in head.items() if re.match(r"(Replaces|Annule|Ersätter|Sostituisce)", k)]
            per = [v for k, v in head.items() if re.match(r"(Service period|Période|Period|Periodo)$", k)][0]
            cur = [v for k, v in head.items() if re.match(r"(Currency|Devise|Valuta)$", k)][0]
            word, year = str(per).rsplit(" ", 1)
            hi = next(i for i, r in enumerate(rows) if r and r[1] is not None and str(r[1]).startswith("KS")) - 1
            lines = []
            for r in rows[hi + 1:]:
                if not r or r[0] is None:
                    break
                lines.append(dict(case=r[1], amount=float(r[5]) + float(r[6]), cur=cur, y=int(year),
                                  m=MONTHS[word.lower()]))
            sheets.append((number, repl[0] if repl else None, lines))
        gone = {r for _, r, _ in sheets if r}
        for number, _, lines in sheets:
            if number not in gone:
                out += lines
    return pd.DataFrame(out)


def asks(t):
    cases = pd.read_csv(os.path.join(t, "service_cases_2024-09_2026-09.csv"), dtype=str).fillna("")
    rep = cases[(cases.fault_code == "CAP") & (cases.warranty == "WTY") & (cases.action == "REPLACED")]
    inv = invoice_lines(t).drop_duplicates("case")
    p = rep[rep.region != "DACH"].merge(inv, left_on="case_no", right_on="case")
    p = p[((p.y == 2025) & (p.m >= 10)) | ((p.y == 2026) & (p.m <= 9))].copy()
    fx = pd.read_csv(os.path.join(t, "treasury_month_end_rates_2025-2026.csv"))
    fx["y"] = fx.month_end.str[:4].astype(int)
    fx["m"] = fx.month_end.str[5:7].astype(int)
    p = p.merge(fx.rename(columns={"currency": "cur"})[["y", "m", "cur", "units_per_eur"]], on=["y", "m", "cur"],
                how="left")
    p["rate"] = p.units_per_eur.fillna(1.0)
    car = pd.read_csv(os.path.join(t, "carrier_statement_inbound_returns_2025-09_2026-10.csv"), dtype={"consignment": str})
    car["share"] = car.charge_eur / car.pieces
    ret = pd.read_excel(os.path.join(t, "depot_returns_log_2024-09_2026-09.xlsx"), dtype={"Consignment": str,
                                                                                          "Reference": str})
    p = p.merge(ret[["Reference", "Consignment"]], left_on="case_no", right_on="Reference", how="left")
    p = p.merge(car[["consignment", "share"]], left_on="Consignment", right_on="consignment", how="left")
    p["pack"] = np.where(p.y == 2025, 52.40, 49.80)
    cost = pd.read_excel(os.path.join(t, "service_cost_sheet_fy2026.xlsx"), sheet_name="Service pack", header=3)
    k = cost[cost.Part == "4100-8820"]["Standard cost EUR"].tolist()
    assert k == [52.40, 49.80], k
    p["cost"] = p.amount / p.rate + p.share.fillna(0.0) + p.pack
    a = p.groupby("partner").cost.agg(["count", "mean"])
    # the months: the depot from SAP and its own cases, the partners from their invoices
    sap = pd.read_excel(os.path.join(t, "depot_service_orders_2024-09_2026-03.xlsx"), dtype={"Auftrag": str})
    s = sap[(sap.Schadenscode == "B10") & (sap.Auftragsart == "ZS01") & sap["Techn. Abschluss"].notna() &
            sap["Einbau-SN"].notna()]
    sd = pd.to_datetime(s["Techn. Abschluss"], format="%d.%m.%Y")
    dach = list(zip(sd.dt.year, sd.dt.month))
    dc = rep[(rep.region == "DACH") & (rep.closed_utc != "") & ~rep.legacy_ref.isin(set(s.Auftrag))]
    dd = pd.to_datetime(dc.closed_utc)
    dach += list(zip(dd.dt.year, dd.dt.month))
    allp = rep[rep.region != "DACH"].merge(inv, left_on="case_no", right_on="case")
    months = pd.Series(dach + list(zip(allp.y, allp.m))).value_counts()
    win = [(2025, m) for m in (10, 11, 12)] + [(2026, m) for m in range(1, 10)]
    return a, {ym: int(months.get(ym, 0)) for ym in win}


def main(target, golden):
    g, n, s2t, s2c = reserve(target)
    wb = openpyxl.load_workbook(os.path.join(golden, "z2638_warranty_reserve.xlsx"), data_only=True)
    book = wb["Reserve"]["B5"].value
    rb = int(math.floor(g / 50 + 0.5) * 50)
    ok = [("reserve", rb == book and n == 7200, "%.2f packs (%d packs, tray %.4f, cell %.4f) -> %d; workbook %s"
           % (g, n, s2t, s2c, rb, book))]
    a, months = asks(target)
    got = {r[0]: r for r in wb["Partner cost"].iter_rows(min_row=5, max_row=8, values_only=True)}
    names = {"DVG": "Dauvergne Réparation (FRB)", "TMR": "Tolley Marsh Repair (UKI)", "LJQ": "Ljungqvist Service (NOR)",
             "FRC": "Ferrucci Assistenza (SOU)"}
    for p_, nm in names.items():
        cnt, avg = int(a.loc[p_, "count"]), float(a.loc[p_, "mean"])
        ok.append(("ask A %s" % p_, got[nm][1] == cnt and got[nm][5] == int(math.floor(avg + 0.5)),
                   "%d at %.4f; workbook %s, %s" % (cnt, avg, got[nm][1], got[nm][5])))
    wm = [r[6] for r in wb["Monthly replacements"].iter_rows(min_row=5, max_row=16, values_only=True)]
    ok.append(("ask B", wm == list(months.values()), "%s; workbook %s" % (list(months.values()), wm)))
    bad = [o for o in ok if not o[1]]
    for name, good, det in ok:
        print("%s  %-10s %s" % ("ok  " if good else "FAIL", name, det))
    print("verify.py: %d figures recomputed from the bundle, %d failed" % (len(ok), len(bad)))
    if bad:
        raise SystemExit(1)


if __name__ == "__main__":
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "target"),
         sys.argv[2] if len(sys.argv) > 2 else os.path.join(here, "golden"))
