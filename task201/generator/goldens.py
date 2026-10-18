#!/usr/bin/env python3
"""task201: the two deliverables and generator/figures.json, computed from the cut pack's frames."""
import os, sys, json, datetime
import numpy as np, pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
TASK = os.path.dirname(HERE)
GOLD = os.path.join(TASK, "golden")
import params as P
import drill

# 2024 is filed in hours, 2025 in minutes (each stated in its own table header); params holds the
# one copy the prior-year report also prints
RET24, SAIFI24, RET25, SAIFI25 = P.RET24, P.SAIFI24, P.RET25, P.SAIFI25


def main_call(st):
    """The drill on the golden basis, and every box's records-fixed counts."""
    ev = st["ev"]
    cl = ev[(ev.rodzaj == "awaryjne") & ev.t_zalaczenie.notna()].reset_index(drop=True)
    e, info = drill.rungs(cl)
    path, boxes = drill.walk(e, "ch_r3")
    p1, boxes1 = drill.walk(e, "ch_r1")
    lv = ["oddzial", "rejon", "klasa_odcinka"]
    counts = []
    scope = e
    for i, l in enumerate(lv):
        g = scope.groupby(l).agg(przerwy=("zdarzenie_id", "size"),
                                 odbiorcoprzerwy=("odbiorcy", "sum"),
                                 godziny=("ch_r3", "sum"))
        order = boxes[i]["shares"].index
        counts.append(g.reindex(order))
        scope = scope[scope[l] == path[i]]
    out = dict(
        path=path, naive_path=p1,
        branch=path[0], district=path[1], klass=path[2],
        leaf_share_pct=int(round(float(boxes[2]["shares"].iloc[0]))),
        leaf_share_exact=float(boxes[2]["shares"].iloc[0]),
        leaf_odbiorcoprzerwy=int(counts[2].odbiorcoprzerwy[path[2]]),
        runner_branch=str(boxes[0]["shares"].index[1]),
        runner_klass=str(boxes[2]["shares"].index[1]),
        levels=[{ "level": lv[i],
                  "boxes": [{"name": str(k), "przerwy": int(counts[i].przerwy[k]),
                             "odbiorcoprzerwy": int(counts[i].odbiorcoprzerwy[k]),
                             "udzial_pct": int(round(float(boxes[i]["shares"][k]))),
                             "udzial_exact": float(boxes[i]["shares"][k]),
                             "godziny": float(counts[i].godziny[k])}
                            for k in counts[i].index]} for i in range(3)],
        delay_by_class={k: float(v) for k, v in info["by_cls"].items()},
        pooled_delay=float(info["pooled"]),
        rungs={c: {k: float(v) for k, v in e.groupby("oddzial")[c].sum().items()}
               for c in ["ch_r1", "ch_r2", "ch_r3"]},
        n_events=int(len(cl)), n_telemetered=int(cl.telemetrowane.sum()))
    return out, e, boxes


def ask_a1(st):
    """Planned switching: customer-hours per branch and month, and the scheduled JOBS that overran.

    The grafik holds one row per job and the dispatch export one row per interruption, so a job that
    switched out two or three sections of one feeder in a visit appears once in the schedule and
    several times in the export. The window is the job's, so the overrun is the job's too.
    """
    pl = st["pl"].copy()
    pl["mies"] = pd.to_datetime(pl.t_utrata).dt.strftime("%Y-%m")
    pl["godziny"] = pl.odbiorcy * (pl.t_zalaczenie - pl.t_utrata).dt.total_seconds() / 3600.0
    ch = pl.groupby(["oddzial", "mies"]).godziny.sum().round(1)
    jobs = pl.groupby("id_pracy").agg(oddzial=("oddzial", "first"),
                                      okno_do=("okno_do", "max"),
                                      zalaczenie=("t_zalaczenie", "max"))
    jobs["over_min"] = ((jobs.zalaczenie - jobs.okno_do).dt.total_seconds() / 60.0).clip(lower=0)
    over = jobs[jobs.over_min > 0]
    per_branch = over.groupby("oddzial").size()
    rows_over = pl[pl.przekroczenie_min > 0]
    ohours = pl[pl.id_pracy.isin(set(over.index))].groupby("oddzial").godziny.sum().round(1)
    nhours = pl[pl.przekroczenie_min > 0].groupby("oddzial").godziny.sum().round(1)
    return dict(
        godziny={"%s|%s" % k: float(v) for k, v in ch.items()},
        prace_przekroczone={b: int(per_branch.get(b, 0)) for b in P.BRANCH},
        prac_przekroczonych_total=int(len(over)),
        godziny_prac_przekroczonych={b: float(ohours.get(b, 0.0)) for b in P.BRANCH},
        naive_godziny_wierszy={b: float(nhours.get(b, 0.0)) for b in P.BRANCH},
        przekroczenie_godz_total=float(round(over.over_min.sum() / 60.0, 1)),
        prac_planowanych=int(len(jobs)),
        wierszy_planowanych=int(len(pl)),
        naive_rows_over={b: int(rows_over.groupby("oddzial").size().get(b, 0)) for b in P.BRANCH},
        naive_rows_over_godz=float(round(rows_over.przekroczenie_min.sum() / 60.0, 1)))


def ask_a2(st):
    """Notifications no interruption record matches, by branch and channel, computed on the shipped
    log: the agent re-save rows are one notification, and the rows the IVR left without a device
    resolve through the customer register."""
    nl = st["nl_shipped"].copy()
    dv = st["dv"][["urzadzenie_id", "oddzial", "odbiorcy_za_urzadzeniem"]]
    ppe = st["ppe"][["nr_ppe", "id_urzadzenia"]]
    nl["ts"] = pd.to_datetime(nl.t_przyjecia)
    nld = nl.drop_duplicates(["nr_odbiorcy", "t_przyjecia", "opis", "kanal"])
    orph = nld[(nld.zdarzenie_id == "") | nld.zdarzenie_id.isna()].copy()
    pm = dict(zip(ppe.nr_ppe, ppe.id_urzadzenia))
    orph["dev"] = orph.id_urzadzenia.replace("", np.nan).fillna(orph.nr_odbiorcy.map(pm))
    br = dict(zip(dv.urzadzenie_id, dv.oddzial)); cu = dict(zip(dv.urzadzenie_id, dv.odbiorcy_za_urzadzeniem))
    orph["br"] = orph.dev.map(br); orph["odb"] = orph.dev.map(cu)
    g = orph.groupby(["br", "kanal"]).size()
    return dict(
        bez_zdarzenia={"%s|%s" % k: int(v) for k, v in g.items()},
        bez_zdarzenia_total=int(len(orph)),
        odbiorcy_za_urzadzeniami={b: int(orph[orph.br == b].odb.sum()) for b in P.BRANCH},
        powtorzenia_crm=int(len(nl) - len(nld)),
        naive_total=int(((nl.zdarzenie_id == "") | nl.zdarzenie_id.isna()).sum()),
        nierozwiazane_bez_rejestru=int(orph.id_urzadzenia.replace("", np.nan).isna().sum()))


def ask_a3(st):
    """The October 2026 revision to reported unplanned interruption minutes, per voltage level."""
    ret, prelim = st["returns"], st["prelim"]
    out = dict()
    # Every row of the preliminary sheet carries the earlier customer count, so the two sheets are
    # re-based on one denominator before the October difference is read. The ratio is the one the
    # untouched months all share.
    m = ret.miesiac != "2026-10"
    ratio = float((prelim[m].saidi_planowane_min_na_odb.sum() / ret[m].saidi_planowane_min_na_odb.sum()))
    rev = {}
    for v in ["SN", "nN"]:
        a = float(ret[(ret.miesiac == "2026-10") & (ret.poziom == v)].saidi_awaryjne_min_na_odb.iloc[0])
        b = float(prelim[(prelim.miesiac == "2026-10") & (prelim.poziom == v)].saidi_awaryjne_min_na_odb.iloc[0]) / ratio
        rev["%s|min" % v] = round(a - b, 1)
        rev["%s|pct" % v] = round((a - b) / b * 100.0, 1)
        rev["%s|min_naive" % v] = round(a - b * ratio, 1)
        rev["%s|pct_naive" % v] = round((a - b * ratio) / (b * ratio) * 100.0, 1)
    rev["denominator_ratio"] = round(ratio, 5)
    out["korekta_pazdziernik"] = rev
    return out


def write_note(call, a1, a2, a3, path):
    from docx import Document
    from docx.shared import Pt, Mm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    nm = {b: P.BRANCH[b]["name"] for b in P.BRANCH}
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Mm(22); s.top_margin = s.bottom_margin = Mm(18)
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)

    h = d.add_paragraph()
    r = h.add_run("Slot remediacyjny 2026/2027 idzie na %s w rejonie %s"
                  % (P.CLS_PL[call["klass"]], call["district"]))
    r.bold = True; r.font.size = Pt(14)
    p = d.add_paragraph("Nota PR-07/2026-11/GRB. Do: narada ruchowa oddzialu %s. Od: Biuro Analiz "
                        "Sieciowych. PR-07 krok 5. Okno analizy 1 pazdziernika do 15 listopada 2026, "
                        "dane pobrane 16 listopada 2026 o 06:00." % nm[call["branch"]])
    p.runs[0].font.size = Pt(8.5); p.runs[0].font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    d.add_paragraph()
    pr = d.add_paragraph()
    pr.add_run("Przypisanie. ").bold = True
    pr.add_run("Slot remediacyjny przypada na klase odcinkow %s w rejonie %s, oddzial %s (%s). "
               "Wskazana klasa odpowiada za %s odbiorco-przerw w swoim rejonie i za %d procent "
               "odbiorco-godzin przerw rejonu. Oddzial drugi w kolejnosci na poziomie pierwszym "
               "to %s." % (P.CLS_PL[call["klass"]], call["district"], nm[call["branch"]],
               call["branch"], "{:,}".format(call["leaf_odbiorcoprzerwy"]).replace(",", " "),
               call["leaf_share_pct"], nm[call["runner_branch"]]))

    d.add_paragraph()
    t = d.add_paragraph(); t.add_run("Poziom 1, oddzialy").bold = True
    tb = d.add_table(rows=1, cols=4); tb.style = "Light Grid Accent 1"
    for i, c in enumerate(["Oddzial", "Udzial w sieci", "Przerwy awaryjne", "Odbiorco-przerwy"]):
        tb.rows[0].cells[i].text = c
    for b in call["levels"][0]["boxes"]:
        row = tb.add_row().cells
        row[0].text = "%s (%s)" % (nm[b["name"]], b["name"])
        row[1].text = "%d %%" % b["udzial_pct"]
        row[2].text = "{:,}".format(b["przerwy"]).replace(",", " ")
        row[3].text = "{:,}".format(b["odbiorcoprzerwy"]).replace(",", " ")


    sp = d.add_paragraph("Zrodlo: SD-OMS, eksporty wylaczen 2026-10-01..2026-11-15; rejestr urzadzen SN i nN, stan 2026-10-01; archiwum telemetrii. Przerwy awaryjne wg PR-07 s.3.")
    sp.runs[0].font.size = Pt(7.6); sp.runs[0].font.color.rgb = RGBColor(0x70, 0x70, 0x70)
    d.add_paragraph()
    t = d.add_paragraph(); t.add_run("Poziom 2, rejony oddzialu %s" % nm[call["branch"]]).bold = True
    tb = d.add_table(rows=1, cols=4); tb.style = "Light Grid Accent 1"
    for i, c in enumerate(["Rejon", "Udzial w oddziale", "Przerwy awaryjne", "Odbiorco-przerwy"]):
        tb.rows[0].cells[i].text = c
    for b in call["levels"][1]["boxes"]:
        row = tb.add_row().cells
        row[0].text = b["name"]; row[1].text = "%d %%" % b["udzial_pct"]
        row[2].text = "{:,}".format(b["przerwy"]).replace(",", " ")
        row[3].text = "{:,}".format(b["odbiorcoprzerwy"]).replace(",", " ")

    d.add_paragraph()
    t = d.add_paragraph(); t.add_run("Poziom 3, klasy odcinkow rejonu %s" % call["district"]).bold = True
    tb = d.add_table(rows=1, cols=4); tb.style = "Light Grid Accent 1"
    for i, c in enumerate(["Klasa odcinka", "Udzial w rejonie", "Przerwy awaryjne", "Odbiorco-przerwy"]):
        tb.rows[0].cells[i].text = c
    for b in call["levels"][2]["boxes"]:
        row = tb.add_row().cells
        row[0].text = b["name"]; row[1].text = "%d %%" % b["udzial_pct"]
        row[2].text = "{:,}".format(b["przerwy"]).replace(",", " ")
        row[3].text = "{:,}".format(b["odbiorcoprzerwy"]).replace(",", " ")

    d.add_page_break()
    t = d.add_paragraph(); t.add_run("Zalacznik A. Wylaczenia planowane w oknie analizy").bold = True
    tb = d.add_table(rows=1, cols=4); tb.style = "Light Grid Accent 1"
    for i, c in enumerate(["Oddzial", "Odbiorco-godziny prac po slocie",
                           "Odbiorco-godziny prac planowanych 2026-10",
                           "Odbiorco-godziny prac planowanych 2026-11"]):
        tb.rows[0].cells[i].text = c
    for b in P.BRANCH:
        row = tb.add_row().cells
        row[0].text = "%s (%s)" % (nm[b], b)
        row[1].text = "%.1f" % a1["godziny_prac_przekroczonych"][b]
        row[2].text = "%.1f" % a1["godziny"].get("%s|2026-10" % b, 0.0)
        row[3].text = "%.1f" % a1["godziny"].get("%s|2026-11" % b, 0.0)
    d.add_paragraph("Prac planowanych w oknie analizy: %d, na %d wylaczeniach. Poza opublikowane "
                    "okno wyszlo %d prac, lacznie o %.1f godziny."
                    % (a1["prac_planowanych"], a1["wierszy_planowanych"],
                       a1["prac_przekroczonych_total"], a1["przekroczenie_godz_total"]))

    d.add_paragraph()
    t = d.add_paragraph(); t.add_run("Zalacznik B. Zgloszenia bez powiazanego wylaczenia").bold = True
    tb = d.add_table(rows=1, cols=5); tb.style = "Light Grid Accent 1"
    for i, c in enumerate(["Oddzial", "telefon", "IVR", "aplikacja", "Odbiorcy za urzadzeniami"]):
        tb.rows[0].cells[i].text = c
    for b in P.BRANCH:
        row = tb.add_row().cells
        row[0].text = "%s (%s)" % (nm[b], b)
        for j, k in enumerate(["telefon", "IVR", "aplikacja"]):
            row[1 + j].text = str(a2["bez_zdarzenia"].get("%s|%s" % (b, k), 0))
        row[4].text = "{:,}".format(a2["odbiorcy_za_urzadzeniami"][b]).replace(",", " ")
    d.add_paragraph("Razem zgloszen bez powiazanego wylaczenia: %d." % a2["bez_zdarzenia_total"])

    d.add_paragraph()
    t = d.add_paragraph(); t.add_run("Zalacznik C. Wskazniki sprawozdawcze koncesji").bold = True
    tb = d.add_table(rows=1, cols=5); tb.style = "Light Grid Accent 1"
    for i, c in enumerate(["Rok sprawozdawczy", "Poziom napiecia sieci", "SAIDI awaryjne [min]", "SAIDI planowane [min]",
                           "SAIFI awaryjne / planowane"]):
        tb.rows[0].cells[i].text = c
    for y, tabv in [("2024", RET24), ("2025", RET25)]:
        for v in ["SN", "nN"]:
            row = tb.add_row().cells
            row[0].text = y; row[1].text = v
            row[2].text = "%.1f" % (tabv[(v, "awaryjne")] * (60.0 if y == "2024" else 1.0))
            row[3].text = "%.1f" % (tabv[(v, "planowane")] * (60.0 if y == "2024" else 1.0))
            sf = SAIFI24 if y == "2024" else SAIFI25
            row[4].text = "%.2f / %.2f" % (sf[(v, "awaryjne")], sf[(v, "planowane")])
    d.add_paragraph("Korekta pazdziernika 2026 do minut awaryjnych: SN %+.1f min (%+.1f procent), "
                    "nN %+.1f min (%+.1f procent)."
                    % (a3["korekta_pazdziernik"]["SN|min"], a3["korekta_pazdziernik"]["SN|pct"],
                       a3["korekta_pazdziernik"]["nN|min"], a3["korekta_pazdziernik"]["nN|pct"]))

    nt = d.add_paragraph(
        "Uwagi do liczb. Identyfikatory urzadzen z eksportow SD-OMS sprowadzono do postaci z zerami "
        "wiodacymi, zgodnie ze slownikiem danych; bez tego 2 823 wiersze eksportu nie lacza sie z "
        "rejestrem. Piec przerw bez zapisanego zalaczenia na 16 listopada nie wchodzi do pomiaru, "
        "zgodnie z PR-07 s.3. Grafik wylaczen planowanych czytany wg wydania 2 dla listopada, ktore "
        "zastepuje wydanie 1. Udzialy podano w pelnych procentach, liczone z niezaokraglonych "
        "odbiorco-godzin.")
    nt.runs[0].font.size = Pt(7.8); nt.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    pf = d.add_paragraph("Biuro Analiz Sieciowych BEO | nota PR-07/2026-11/GRB | 16 listopada 2026")
    pf.runs[0].font.size = Pt(8); pf.runs[0].font.color.rgb = RGBColor(0x70, 0x70, 0x70)
    cp = d.core_properties
    cp.author = "Biuro Analiz Sieciowych BEO"; cp.title = "Nota PR-07/2026-11/GRB, slot remediacyjny 2026/2027"
    cp.comments = ""; cp.last_modified_by = "Biuro Analiz Sieciowych BEO"
    cp.created = datetime.datetime(2026, 11, 16, 7, 40)
    cp.modified = datetime.datetime(2026, 11, 16, 7, 40)
    cp.revision = 3
    d.save(path)


def write_tree(call, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch
    nm = {b: P.BRANCH[b]["name"] for b in P.BRANCH}
    fig, ax = plt.subplots(figsize=(12.4, 7.0), dpi=150)
    ax.set_xlim(0, 100); ax.set_ylim(0, 100); ax.axis("off")
    ON, OFF, EDGE = "#1f4e79", "#dfe4ea", "#5b6770"

    def box(x, y, w, h, label, pct, extra, lit):
        p = FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.6,rounding_size=1.2",
                           linewidth=1.9 if lit else 0.9,
                           edgecolor=ON if lit else EDGE, facecolor=ON if lit else OFF)
        ax.add_patch(p)
        c = "white" if lit else "#1a1a1a"
        ax.text(x + w / 2, y + h * 0.66, label, ha="center", va="center", fontsize=9.2,
                color=c, fontweight="bold" if lit else "normal")
        ax.text(x + w / 2, y + h * 0.30, "%s  |  %d %% odb.-godzin" % (extra, pct),
                ha="center", va="center", fontsize=7.3, color=c)

    ax.text(50, 96, "Ekspozycja sezonu siedzi na odgalezieniach SN w rejonie %s" % call["district"],
            ha="center", fontsize=12.5, fontweight="bold", color="#1a1a1a")
    ax.text(50, 92, "Drzewo PR-07 krok 1 do 3; w kazdym polu odbiorco-przerwy i udzial poziomu; "
                    "okno 1.10 do 15.11.2026", ha="center", fontsize=9, color="#5b6770")
    rows = [("Poziom 1, oddzialy", 70.0, call["levels"][0], call["branch"], lambda k: nm.get(k, k)),
            ("Poziom 2, rejony %s" % nm[call["branch"]], 42.0, call["levels"][1], call["district"], lambda k: k),
            ("Poziom 3, klasy odcinkow %s" % call["district"], 14.0, call["levels"][2], call["klass"], lambda k: k)]
    centres = {}
    for title, y, lvl, lit_name, labf in rows:
        ax.text(98, y + 15.3, title, fontsize=9.2, fontweight="bold", color="#1f4e79",
                ha="right")
        n = len(lvl["boxes"]); w = 21.0; gapw = (100 - 4 - n * w) / max(1, n - 1)
        for i, b in enumerate(lvl["boxes"]):
            x = 2 + i * (w + gapw)
            box(x, y, w, 13.0, labf(b["name"]), b["udzial_pct"],
                "%s odb.-przerw" % "{:,}".format(b["odbiorcoprzerwy"]).replace(",", " "),
                b["name"] == lit_name)
            if b["name"] == lit_name:
                centres[title] = (x + w / 2, y)
    ks = list(centres)
    for a, b in zip(ks, ks[1:]):
        x0, y0 = centres[a]; x1, y1 = centres[b]
        ax.annotate("", xy=(x1, y1 + 13.2), xytext=(x0, y0),
                    arrowprops=dict(arrowstyle="-|>", color=ON, linewidth=2.1,
                                    connectionstyle="arc3,rad=0.0"))
    ax.text(50, 6.0, "Slot remediacyjny: %s, rejon %s, oddzial %s. Wskazana klasa: %s "
                     "odbiorco-przerw w rejonie." % (P.CLS_PL[call["klass"]], call["district"],
                     nm[call["branch"]], "{:,}".format(call["leaf_odbiorcoprzerwy"]).replace(",", " ")),
            ha="center", fontsize=10.2, fontweight="bold", color="#1f4e79")
    ax.text(50, 2.4, "Biuro Analiz Sieciowych BEO | 16 listopada 2026 | zrodlo: SD-OMS, rejestr "
                     "urzadzen SN i nN, archiwum telemetrii, zgloszenia klientow",
            ha="center", fontsize=7.4, color="#6b7280")
    fig.savefig(path, bbox_inches="tight", facecolor="white", metadata={"Software": None})
    plt.close(fig)


def write_all(st):
    os.makedirs(GOLD, exist_ok=True)
    call, e, boxes = main_call(st)
    a1, a2, a3 = ask_a1(st), ask_a2(st), ask_a3(st)
    write_note(call, a1, a2, a3, os.path.join(GOLD, "plan_remediacji_2026.docx"))
    write_tree(call, os.path.join(GOLD, "drzewo_udzialow_2026.png"))
    # every figure the write-up states, emitted rather than typed
    ev = st["ev"]
    cl = ev[(ev.rodzaj == "awaryjne") & ev.t_zalaczenie.notna()]
    e2, info2 = drill.rungs(cl.reset_index(drop=True))
    hrs = (e2.groupby("oddzial").ch_r3.sum() / 1000).round(1)
    s = hrs.sort_values(ascending=False)
    ov = cl[cl.telemetrowane]
    tel_share = {b: int(round(100 * cl[cl.oddzial == b].telemetrowane.mean())) for b in P.BRANCH}
    trunk_station = int(round(100 * ov.klasa_odcinka.isin(["SN-magistrala", "SN/nN-stacja"]).mean()))
    narrative = dict(
        export_rows=int(st["counts"]["legacy"] + st["counts"]["current"]),
        interruptions_all=int(len(ev)),
        raw_join_lost=int(st["raw_join_lost"]),
        population=int(len(cl)), timed=int(cl.telemetrowane.sum()),
        untimed=int((~cl.telemetrowane).sum()),
        hours_thousand={k: float(v) for k, v in hrs.items()},
        lead=float(round(s.iloc[0] / s.iloc[1], 2)),
        telemetered_share_pct=tel_share,
        overlap_trunk_station_pct=trunk_station)
    figs = dict(call=call, a1=a1, a2=a2, a3=a3, counts=st["counts"], narrative=narrative)
    json.dump(figs, open(os.path.join(HERE, "figures.json"), "w"), indent=1, sort_keys=True,
              default=float)
    return figs
