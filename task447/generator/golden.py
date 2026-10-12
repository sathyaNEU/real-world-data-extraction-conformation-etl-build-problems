"""Every graded figure, computed from the shipped files under target/ alone, the way rule OHCCR 4 and
the market oversight appendix read: no generator state, no seeds.

    compute(target) -> dict of figures and the intermediate tables the write-up and the checks use
"""
import csv
import json
import os
from collections import defaultdict
from datetime import date, timedelta

import openpyxl
import pandas as pd

RUNOUT = pd.Timestamp("2025-06-30")
BASE_START, BASE_END = pd.Timestamp("2024-01-01"), pd.Timestamp("2024-12-31")
THRESHOLD = 2_000_000
LINES = {"IMG": "imaging", "ECHO": "echocardiography", "ENDO": "endoscopy", "ARTH": "arthroscopy",
         "LAB": "laboratory"}


def family(cpt):
    """Appendix A1 service lines."""
    try:
        c = int(cpt)
    except ValueError:
        return None
    if 70010 <= c <= 76999 or 77046 <= c <= 77067:
        return "IMG"
    if 93303 <= c <= 93356:
        return "ECHO"
    if 43235 <= c <= 43270 or 45378 <= c <= 45398:
        return "ENDO"
    if 29805 <= c <= 29999:
        return "ARTH"
    if 80047 <= c <= 89398:
        return "LAB"
    return None


# ------------------------------------------------------------------ loading and conforming
def load(target):
    p = lambda f: os.path.join(target, f)
    fac24 = pd.read_csv(p("apcd_op_facility_claims_2024.csv"), dtype=str, keep_default_na=False)
    v2 = pd.read_csv(p("apcd_op_facility_claims_2021_2023.csv"), dtype=str, keep_default_na=False)
    v2 = v2.rename(columns={
        "SUBMITTER_ID": "submitter_code", "CLAIM_ID": "payer_claim_control_number", "LINE_NO": "line_counter",
        "CLM_STATUS": "claim_frequency_code", "ORIG_CLAIM_ID": "original_claim_control_number", "FORM": "form_type",
        "MEMBER_ID": "member_id", "PRODUCT_ID": "product_id", "DOS_FROM": "service_from_date", "PAID_DT": "paid_date",
        "TOB": "type_of_bill", "POS": "place_of_service", "REV_CD": "revenue_code", "PROC_CD": "procedure_code",
        "MOD1": "procedure_modifier_1", "UNITS": "quantity", "ALLOWED_AMT": "allowed_amount",
        "PLAN_PAID_AMT": "plan_paid_amount", "MBR_RESP_AMT": "member_liability_amount",
        "BILL_NPI": "billing_provider_npi", "BILL_TIN": "billing_provider_tin", "FAC_NPI": "service_facility_npi",
        "REF_NPI": "referring_provider_npi"})
    v2["claim_frequency_code"] = v2["claim_frequency_code"].map({"O": "1", "R": "7", "V": "8"})
    v2["service_from_date"] = pd.to_datetime(v2["service_from_date"], format="%Y%m%d").dt.strftime("%Y-%m-%d")
    v2["paid_date"] = pd.to_datetime(v2["paid_date"], format="%Y%m%d").dt.strftime("%Y-%m-%d")
    v2["billing_provider_tin"] = v2["billing_provider_tin"].str[:2] + "-" + v2["billing_provider_tin"].str[2:]
    v2["service_facility_zip"] = ""
    fac = pd.concat([v2[fac24.columns], fac24], ignore_index=True)
    prof = pd.read_csv(p("apcd_professional_claims_2024.csv"), dtype=str, keep_default_na=False)
    for df in (fac, prof):
        df["dos"] = pd.to_datetime(df["service_from_date"])
        df["paid"] = pd.to_datetime(df["paid_date"])
        df["qty"] = df["quantity"].astype(int)
        df["allowed"] = df["allowed_amount"].astype(float)
    elig = pd.read_csv(p("apcd_eligibility_spans.csv"), dtype=str, keep_default_na=False)
    elig["start"] = pd.to_datetime(elig["coverage_start"])
    elig["end"] = pd.to_datetime(elig["coverage_end"])
    elig["plan_year"] = pd.to_datetime(elig["benefit_period_begin"]).dt.year
    prodref = pd.read_csv(p("apcd_product_reference.csv"), dtype=str, keep_default_na=False)
    tiers = json.load(open(p("network_tier_placements_2020_2026.json")))
    register = pd.read_csv(p("provider_ownership_register.csv"), dtype=str, keep_default_na=False)
    reassign = pd.read_csv(p("npi_reassignment_history.csv"), dtype=str, keep_default_na=False)
    provfile = pd.read_csv(p("carrier_provider_file.csv"), dtype=str, keep_default_na=False)
    facdir = pd.read_csv(p("facility_directory.csv"), dtype=str, keep_default_na=False)
    medicare = pd.read_csv(p("cms_opps_clfs_2025_wisconsin.csv"), dtype=str, keep_default_na=False)
    wb = openpyxl.load_workbook(p("N25-014_exhibit_C_ledgestone_rates.xlsx"), data_only=True)
    return dict(fac=fac, prof=prof, elig=elig, prodref=prodref, tiers=tiers, register=register, reassign=reassign,
                provfile=provfile, facdir=facdir, medicare=medicare, wb=wb)


def latest_versions(df):
    """Rule 4.02(c) and the dictionary: a replacement supersedes the claim it names, a void cancels it."""
    named = set(df.loc[df["claim_frequency_code"].isin(["7", "8"]), "original_claim_control_number"])
    keep = df["claim_frequency_code"].isin(["1", "7"]) & ~df["payer_claim_control_number"].isin(named)
    return df[keep].copy()


def exhibit_rates(wb, medicare, prodref):
    """Ledgestone's filed rates: facility by (carrier name, cpt), professional by (carrier name, cpt, modifier)."""
    car_name = dict(zip(prodref["submitter_code"], prodref["carrier_name"]))
    fac = {}
    ws = wb["Outpatient facility"]
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r and r[0] == "CPT/HCPCS")
    hdr = rows[hdr_i]
    for r in rows[hdr_i + 1:]:
        if not r or r[0] is None:
            continue
        for j in range(2, len(hdr)):
            if hdr[j]:
                fac[(hdr[j], str(r[0]))] = float(r[j])
    mc = dict(zip(medicare["hcpcs"], medicare["cy2025_amount_wi_cbsa24580"].astype(float)))
    ws = wb["Bayfield terms"]
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r and r[0] == "Service category")
    for r in rows[hdr_i + 1:]:
        if not r or r[0] is None:
            continue
        for cpt in str(r[1]).split(", "):
            fac[("Bayfield Mutual Health", cpt)] = round(mc[cpt] * float(r[2]), 2)
    prof = {}
    ws = wb["Professional"]
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r and r[0] == "CPT/HCPCS")
    hdr = rows[hdr_i]
    for r in rows[hdr_i + 1:]:
        if not r or r[0] is None:
            continue
        for j in range(3, len(hdr)):
            prof[(hdr[j], str(r[0]), r[1] or "")] = float(r[j])
    return fac, prof, car_name


class Elig:
    def __init__(self, elig):
        self.by = defaultdict(list)
        for m, s, e, prod, py, z in zip(elig["member_id"], elig["start"], elig["end"], elig["product_id"],
                                        elig["plan_year"], elig["member_zip"]):
            self.by[m].append((s, e, prod, py, z))

    def on(self, m, d):
        for s, e, prod, py, z in self.by.get(m, ()):
            if s <= d <= e:
                return prod, py, z
        return None, None, None


class Tiers:
    def __init__(self, tiers, prodref):
        self.tiered = dict(zip(prodref["product_id"], prodref["tiered_network"] == "Y"))
        self.t = {}
        for f in tiers["filings"]:
            for pl in f["hospital_system_placements"]:
                self.t[(f["product_id"], f["plan_year"], pl["hospital_system"])] = pl["tier"]

    def tier(self, product, system, plan_year):
        if system is None or not self.tiered.get(product, False):
            return 1
        return self.t[(product, plan_year, system)]


COUNTIES = ("Brown", "Calumet", "Fond du Lac", "Kewaunee", "Manitowoc", "Outagamie", "Sheboygan", "Winnebago")


def zip_county(target):
    """The data dictionary's ZIP-to-county list (section 6)."""
    import re
    import pdfplumber
    with pdfplumber.open(os.path.join(target, "apcd_extract_data_dictionary.pdf")) as pdf:
        text = " ".join(pg.extract_text() or "" for pg in pdf.pages)
    out = {}
    for c in COUNTIES:
        m = re.search(re.escape(c) + r"\s+((?:\d{5}(?:,\s*)?)+)", text)
        for z in re.findall(r"\d{5}", m.group(1)):
            out[z] = c
    return out


# ------------------------------------------------------------------ the register and billing spans
class Control:
    """Which system owns or controls a TIN on a date, through the register's chains (appendix A4)."""
    QUAL = ("owned", "merged into", "affiliated, controlled")

    def __init__(self, register):
        self.rel = defaultdict(list)    # tin or entity name -> list of (parent name, parent tin, rel, start, end)
        self.tin_name = {}
        for r in register.itertuples():
            q = any(r.relationship.startswith(x) for x in self.QUAL)
            start = pd.Timestamp(r.effective_from)
            end = pd.Timestamp(r.effective_to) if r.effective_to else pd.Timestamp("2099-12-31")
            key = r.entity_tin or r.entity_name
            self.rel[key].append((r.related_entity_name, r.related_entity_tin, q, start, end))
            if r.entity_tin:
                self.tin_name[r.entity_tin] = r.entity_name

    def system_of(self, tin, d, depth=0):
        if depth > 6:
            return None
        for pname, ptin, q, s, e in self.rel.get(tin, ()):
            if q and s <= d <= e:
                if pname in ("Ledgestone Health", "Wolf River Health", "Kinnickinnic Health"):
                    return pname
                up = self.system_of(ptin or pname, d, depth + 1)
                if up:
                    return up
        return None


def billing_spans(reassign):
    spans = defaultdict(list)
    for r in reassign.itertuples():
        s = pd.Timestamp(r.effective_from)
        e = pd.Timestamp(r.effective_to) if r.effective_to else pd.Timestamp("2099-12-31")
        spans[r.npi].append((r.billing_tin, s, e))
    return spans


# ------------------------------------------------------------------ the main call
def compute(target):
    X = load(target)
    fac_all = latest_versions(X["fac"])
    prof_all = latest_versions(X["prof"])
    prodref, facdir = X["prodref"], X["facdir"]
    commercial = set(prodref.loc[prodref["line_of_business"] == "Commercial", "product_id"])
    car_name = dict(zip(prodref["submitter_code"], prodref["carrier_name"]))
    fac_rate, prof_rate, _ = exhibit_rates(X["wb"], X["medicare"], prodref)
    E = Elig(X["elig"])
    T = Tiers(X["tiers"], prodref)
    sys_of_fac = dict(zip(facdir["facility_npi"], facdir["health_system"]))
    spans = billing_spans(X["reassign"])
    ctrl = Control(X["register"])
    zc = zip_county(target)

    # Brightwater's taxpayer number and physicians (notice, reassignment history)
    import re
    import pdfplumber
    with pdfplumber.open(os.path.join(target, "N25-014_notice_of_material_change.pdf")) as pdf:
        text = " ".join(pg.extract_text() or "" for pg in pdf.pages)
    bmg_tin = re.search(r"Brightwater's number\s*\((\d{2}-\d{7})\)", text).group(1)

    def tin_on(npi, d):
        return [t for t, s, e in spans.get(npi, ()) if s <= d <= e]

    bmg_npis = {n for n, lst in spans.items() if any(t == bmg_tin for t, s, e in lst)}

    # Ledgestone's sites and the services it offers (4.02(g)): the codes its sites bill
    ldg_npis = set(facdir.loc[facdir["health_system"] == "Ledgestone Health", "facility_npi"])
    f24 = fac_all[(fac_all["dos"] >= BASE_START) & (fac_all["dos"] <= BASE_END)]
    sysname = lambda r_fac: sys_of_fac.get(r_fac)
    offered = defaultdict(set)
    for npi, cpt in zip(f24["service_facility_npi"], f24["procedure_code"]):
        s = sys_of_fac.get(npi)
        if s and s != "Independent":
            offered[s].add(cpt)

    # ---- baseline (4.03)
    base_f = f24[(f24["paid"] <= RUNOUT) & f24["product_id"].isin(commercial)]
    base_p = prof_all[(prof_all["dos"] >= BASE_START) & (prof_all["dos"] <= BASE_END) & (prof_all["paid"] <= RUNOUT)
                      & prof_all["product_id"].isin(commercial)]
    bp = base_p[base_p["billing_provider_tin"] == bmg_tin]
    prof_rows = []
    for r in bp.itertuples():
        rate = prof_rate[(car_name[r.submitter_code], r.procedure_code, r.procedure_modifier_1)]
        prof_rows.append((car_name[r.submitter_code], r.procedure_code, r.allowed, rate * r.qty))
    P = sum(new - old for _, _, old, new in prof_rows)
    P_base = sum(old for _, _, old, _ in prof_rows)

    is_bmg_ref = [bool(n in bmg_npis and bmg_tin in tin_on(n, d))
                  for n, d in zip(base_f["referring_provider_npi"], base_f["dos"])]
    br = base_f[is_bmg_ref].copy()
    br["system"] = br["service_facility_npi"].map(sys_of_fac)
    br["fam"] = br["procedure_code"].map(family)
    info = [E.on(m, d) for m, d in zip(br["member_id"], br["dos"])]
    br["product"] = [i[0] for i in info]
    br["plan_year"] = [i[1] for i in info]
    br["zip"] = [i[2] for i in info]
    br["county"] = br["zip"].map(zc)
    br["t24"] = [T.tier(p, "Ledgestone Health", py) for p, py in zip(br["product"], br["plan_year"])]
    br["t26"] = [T.tier(p, "Ledgestone Health", 2026) for p in br["product"]]
    br["t25"] = [T.tier(p, "Ledgestone Health", 2025) for p in br["product"]]
    br["carrier"] = br["submitter_code"].map(car_name)
    br["offered"] = br["procedure_code"].isin(offered["Ledgestone Health"])
    br["at_ldg"] = br["system"] == "Ledgestone Health"
    elig_mask = br["offered"] & ~br["at_ldg"]
    br["new_rate"] = [fac_rate.get((c, cpt)) if m else None
                      for c, cpt, m in zip(br["carrier"], br["procedure_code"], elig_mask)]
    br["delta"] = [(nr * q - a) if m else 0.0 for nr, q, a, m in zip(br["new_rate"], br["qty"], br["allowed"], elig_mask)]
    br["eligible"] = elig_mask

    # ---- the history (4.05): acquired practices' referred services before and after each acquisition
    hist = acquisitions(X, fac_all, spans, ctrl, sys_of_fac, offered, E, T, commercial)
    fit = fit_models(hist)
    # 4.05: the selected estimate is stated in whole per cent and the projection uses it as stated
    c1, c2 = whole(fit["tier"]["1"]), whole(fit["tier"]["2"])

    br["share_moving"] = [(c1 if t == 1 else c2) if m else 0.0 for t, m in zip(br["t26"], br["eligible"])]
    referred_inc = float((br["share_moving"] * br["delta"]).sum())
    total = P + referred_inc
    fam_inc = {f: float((br.loc[br["fam"] == f, "share_moving"] * br.loc[br["fam"] == f, "delta"]).sum())
               for f in LINES}
    other_fam = referred_inc - sum(fam_inc.values())

    # ---- the market exhibits (rule 4.07, appendix A2 and A3)
    ex1 = ex1_cells(X, zc, bmg_tin, spans)
    ex2 = ex2_cells(X, zc, bmg_tin, ctrl)

    return dict(total=total, P=P, P_base=P_base, referred=referred_inc, fam=fam_inc, other_fam=other_fam,
                c1=c1, c2=c2, fit=fit, ex1=ex1, ex2=ex2,
                bmg_tin=bmg_tin, bmg_npis=sorted(bmg_npis), br=br, prof_rows=prof_rows, hist=hist,
                call="proceed" if total >= THRESHOLD else "close", threshold=THRESHOLD,
                base_ref=float(br["allowed"].sum()), n_prof=len(prof_rows), n_ref=len(br),
                n_elig=int(br["eligible"].sum()), offered=sorted(offered["Ledgestone Health"]),
                fac_rate=fac_rate, prof_rate=prof_rate)


# ------------------------------------------------------------------ acquisitions and the model fit
def first_system_date(lst, ctrl, system, reg_dates):
    """The first date a physician bills under a number the system owns or controls, and that number."""
    best = None
    for t, s, e in sorted(lst, key=lambda x: x[1]):
        for d in [s] + [x for x in reg_dates if s < x <= e]:
            if ctrl.system_of(t, d) == system:
                if best is None or d < best[0]:
                    best = (d, t)
                break
    return best


def joins(spans, ctrl, register, system, is_phys=lambda n: True):
    """Physicians who joined the system by acquisition (appendix A5), with the date and the practice
    number they came with: {npi: (date, practice tin)}."""
    reg_dates = sorted({pd.Timestamp(x) for x in register["effective_from"]})
    first = {}
    for npi, lst in spans.items():
        if is_phys(npi):
            fs = first_system_date(lst, ctrl, system, reg_dates)
            if fs:
                first[npi] = fs
    starts = defaultdict(set)
    for r in register.itertuples():
        if r.entity_tin and any(r.relationship.startswith(q) for q in Control.QUAL):
            starts[r.entity_tin].add(pd.Timestamp(r.effective_from))
    roster = defaultdict(set)
    for npi, lst in spans.items():
        if is_phys(npi):
            for t, s, e in lst:
                roster[t].add(npi)
    out = {}
    for npi, (d, t) in first.items():
        day_before = d - pd.Timedelta(days=1)
        lst = spans[npi]
        on_t_before = any(x == t and s <= day_before <= e for x, s, e in lst)
        prev = [x for x, s, e in lst if e == day_before or (s <= day_before <= e and x != t)]
        # (a) the number the physician already billed under came under the system's control that day
        if on_t_before and ctrl.system_of(t, day_before) != system:
            out[npi] = (d, t)
            continue
        for p in prev:
            if ctrl.system_of(p, day_before) == system:
                continue
            # (b) the practice the physician billed under came under the system's ownership that day
            if d in starts.get(p, set()) and ctrl.system_of(p, d) == system:
                out[npi] = (d, p)
                break
            # (c) every physician of that practice moved to the system's numbers within 90 days
            members = roster[p]
            moves = [first[m][0] for m in members if m in first]
            if moves and len(moves) == len(members) and (max(moves) - min(moves)).days <= 90:
                left = all(max(e for x, s, e in spans[m] if x == p) < first[m][0] for m in members)
                if left:
                    out[npi] = (d, p)
                    break
    return out


def acquisitions(X, fac_all, spans, ctrl, sys_of_fac, offered, E, T, commercial):
    """Rule 4.05's evidence: referred services of the physicians who joined health systems by acquisition
    (appendix A4 and A5), where the database holds their referred services both before and after the
    acquisition, each service dated against the referring physician's own joining date."""
    prov = X["provfile"]
    tax = dict(zip(prov["provider_npi"], prov["primary_taxonomy"]))
    is_phys = lambda n: tax.get(n, "")[:3] in PHYS_TAX
    joined = {}
    for system in ("Ledgestone Health", "Wolf River Health", "Kinnickinnic Health"):
        for npi, (d, p) in joins(spans, ctrl, X["register"], system, is_phys).items():
            joined[npi] = (d, system, p)
    refs = fac_all[fac_all["referring_provider_npi"].isin(joined.keys()) & fac_all["product_id"].isin(commercial)]
    acq = {}
    for npi, sub in refs.groupby("referring_provider_npi"):
        d = joined[npi][0]
        if (sub["dos"] < d).any() and (sub["dos"] >= d).any():
            acq[npi] = joined[npi]
    h = fac_all[fac_all["referring_provider_npi"].isin(acq.keys())].copy()
    h = h[h["product_id"].isin(commercial) & (h["dos"] <= pd.Timestamp("2024-12-31"))]
    h["acq_date"] = h["referring_provider_npi"].map(lambda n: acq[n][0])
    h["acq_sys"] = h["referring_provider_npi"].map(lambda n: acq[n][1])
    h["prac_tin"] = h["referring_provider_npi"].map(lambda n: acq[n][2])
    h["system"] = h["service_facility_npi"].map(sys_of_fac)
    h["offered"] = [c in offered[s] for c, s in zip(h["procedure_code"], h["acq_sys"])]
    h = h[h["offered"]].copy()
    h["at"] = (h["system"] == h["acq_sys"]).astype(int)
    h["post"] = h["dos"] >= h["acq_date"]
    info = [E.on(m, d) for m, d in zip(h["member_id"], h["dos"])]
    h["product"] = [i[0] for i in info]
    h["plan_year"] = [i[1] for i in info]
    h["tier"] = [T.tier(p, s, py) for p, s, py in zip(h["product"], h["acq_sys"], h["plan_year"])]
    h["tier_cal"] = [T.tier(p, s, d.year) for p, s, d in zip(h["product"], h["acq_sys"], h["dos"])]
    h["carrier"] = h["submitter_code"]
    h["year"] = h["dos"].dt.year
    h["tiered"] = h["product"].map(T.tiered)
    return h


def whole(x):
    """A share stated in whole per cent (rule 4.05)."""
    return round(x, 2) + 0.0


def pre_shares(h):
    pre = h[~h["post"]].groupby("prac_tin")["at"].mean()
    return pre.to_dict()


def fit_models(h, before_year=None, models=None):
    """Capture = moved share of the outside flow, pooled over practices: sum(at - s0) / sum(1 - s0) over
    post-acquisition services, per cell of each candidate model."""
    models = models or MODEL_KEYS
    s0 = pre_shares(h)
    d = h if before_year is None else h[h["year"] < before_year]
    post = d[d["post"]].copy()
    post["s0"] = post["prac_tin"].map(s0)
    post["num"] = post["at"] - post["s0"]
    post["den"] = 1 - post["s0"]
    out = {}
    for name in models:
        keys = pd.Series(model_key(name, post, models), index=post.index)
        g = post.groupby(keys)[["num", "den"]].sum()
        out[name] = (g["num"] / g["den"]).to_dict()
    return out


MODEL_KEYS = {
    "pooled": None,
    "system": lambda df: df["acq_sys"],
    "carrier": lambda df: df["carrier"],
    "product": lambda df: df["product"],
    "product x system": lambda df: ["%s|%s" % (p, s) for p, s in zip(df["product"], df["acq_sys"])],
    "tiered flag": lambda df: df["tiered"],
    "tier": lambda df: df["tier"],
    "tier (calendar year)": lambda df: df["tier_cal"],
}


# finer versions of the tier model a solver could carry to Brightwater's referrals; the records are built so
# every one of them states the same whole-per-cent shares (C1)
REFINEMENTS = {
    "tier x system": lambda df: ["%s|%s" % (t, s) for t, s in zip(df["tier"], df["acq_sys"])],
    "tier x carrier": lambda df: ["%s|%s" % (t, c) for t, c in zip(df["tier"], df["carrier"])],
    "tier x product": lambda df: ["%s|%s" % (t, p) for t, p in zip(df["tier"], df["product"])],
    "tier x service line": lambda df: ["%s|%s" % (t, f) for t, f in zip(df["tier"], df["procedure_code"].map(family))],
    "tier x practice": lambda df: ["%s|%s" % (t, p) for t, p in zip(df["tier"], df["prac_tin"])],
    "tier x year, 2024": lambda df: ["%s|%s" % (t, y) for t, y in zip(df["tier"], df["year"])],
    "tier x plan year, 2024": lambda df: ["%s|%s" % (t, y) for t, y in zip(df["tier"], df["plan_year"])],
}


def model_key(name, df, models=None):
    if name == "pooled":
        return ["all"] * len(df)
    return [str(x) for x in (models or MODEL_KEYS)[name](df)]


def backtest(h, models=None, stated=False):
    """Rule 4.05's criterion: fitted only to years before each full calendar year after each acquisition,
    the largest miss in the share of the practice's referred services that went to its acquiring system.
    stated=True scores each fit as stated in whole per cent instead of as fitted."""
    models = models or MODEL_KEYS
    s0 = pre_shares(h)
    acq_year = h.groupby("prac_tin")["acq_date"].max().dt.year.to_dict()
    res = {}
    for name in models:
        worst, at = 0.0, None
        cells = 0
        for tin, ay in sorted(acq_year.items()):
            for y in range(ay + 1, 2025):
                fit = fit_models(h, before_year=y, models={name: models[name]})[name]
                if stated:
                    fit = {k: whole(v) for k, v in fit.items()}
                test = h[(h["prac_tin"] == tin) & (h["year"] == y)]
                if test.empty:
                    continue
                keys = model_key(name, test, models)
                if any(k not in fit for k in keys):
                    continue
                pred = sum(s0[tin] + fit[k] * (1 - s0[tin]) for k in keys) / len(test)
                obs = test["at"].mean()
                cells += 1
                if abs(pred - obs) > worst:
                    worst, at = abs(pred - obs), (tin, y, obs, pred)
        res[name] = (worst, at, cells)
    return res


PHYS_TAX = ("207", "208")


# ------------------------------------------------------------------ the market exhibits (appendix A2 and A3)
EX1_LINES = ("IMG", "ENDO", "ARTH")
EX_COUNTIES = ("Manitowoc", "Kewaunee")
OFFICE = {str(c) for c in range(99202, 99216)}
SPECIALTY = {"207Q00000X": "primary care", "207R00000X": "primary care", "208000000X": "primary care",
             "207RG0100X": "gastroenterology", "207X00000X": "orthopaedic surgery",
             "207RC0000X": "cardiovascular disease"}
SPECIALTIES = ("primary care", "gastroenterology", "orthopaedic surgery", "cardiovascular disease")


def _versions(df, how):
    if how == "latest":
        return latest_versions(df)
    if how == "as filed":
        return df.copy()
    raise ValueError(how)


def _services(df, key, how):
    """Count services. 'once': a service billed to two plans is one service, so each key counts the units one
    plan holds for it (most units over the plans); 'every claim': the units on every claim line; 'one per key':
    a single service per key whatever the units or plans."""
    if df.empty:
        return 0
    if how == "every claim":
        return int(df["qty"].sum())
    if how == "one per key":
        return int((df.groupby(key)["qty"].sum() > 0).sum())
    per = df.groupby(key + ["submitter_code"])["qty"].sum()
    return int(per.groupby(level=list(range(len(key)))).max().sum())


def ex1_cells(X, zc, bmg_tin, spans, versions="latest", second="once", key=("member_id", "dos", "procedure_code",
              "procedure_modifier_1", "service_facility_npi"), residence="span", runout=False):
    """Appendix A2: the commercial imaging, endoscopy and arthroscopy services each county's residents received in
    2024, and the share whose referring provider billed under Brightwater's number on the date of service."""
    commercial = set(X["prodref"].loc[X["prodref"]["line_of_business"] == "Commercial", "product_id"])
    f = X["fac"]
    f = f[(f["dos"] >= BASE_START) & (f["dos"] <= BASE_END) & f["product_id"].isin(commercial)]
    f = f[f["procedure_code"].map(family).isin(EX1_LINES)]
    if runout:
        f = f[f["paid"] <= RUNOUT]
    f = _versions(f, versions)
    el = X["elig"]
    if residence == "span":
        idx = defaultdict(list)
        for m, p, s_, e_, z in zip(el["member_id"], el["product_id"], el["start"], el["end"], el["member_zip"]):
            idx[(m, p)].append((s_, e_, z))
        zips = [next((z for s_, e_, z in idx.get((m, p), ()) if s_ <= d <= e_), None)
                for m, p, d in zip(f["member_id"], f["product_id"], f["dos"])]
    else:
        first = el.sort_values(["member_id", "start"]).groupby("member_id")["member_zip"]
        pick = (first.first() if residence == "first" else first.last()).to_dict()
        zips = [pick.get(m) for m in f["member_id"]]
    f = f.assign(county=[zc.get(z) for z in zips], line=f["procedure_code"].map(family))
    bmg = [any(t == bmg_tin and s_ <= d <= e_ for t, s_, e_ in spans.get(n, ())) for n, d in
           zip(f["referring_provider_npi"], f["dos"])]
    f = f.assign(bmg=bmg)
    key = list(key)
    out = {}
    for county in EX_COUNTIES:
        for ln in EX1_LINES:
            sel = f[(f["county"] == county) & (f["line"] == ln)]
            tot = _services(sel, key, second)
            b = _services(sel[sel["bmg"]], key, second)
            out[(county, ln)] = (tot, b, 100.0 * b / tot)
    return out


def ex2_cells(X, zc, bmg_tin, ctrl, versions="latest", second="once", specialty="provider", physicians=True,
              system="register", telehealth=False, runout=False):
    """Appendix A3: commercial office visits in the two counties by the rendering physician's specialty, and the
    shares billed under Brightwater's number and under the numbers Ledgestone owns or controls."""
    commercial = set(X["prodref"].loc[X["prodref"]["line_of_business"] == "Commercial", "product_id"])
    p = X["prof"]
    p = p[(p["dos"] >= BASE_START) & (p["dos"] <= BASE_END) & p["product_id"].isin(commercial)]
    p = p[p["procedure_code"].isin(OFFICE) & p["service_facility_zip"].map(zc).isin(EX_COUNTIES)]
    if not telehealth:          # A3: a visit by telehealth is not an office visit
        p = p[(p["place_of_service"] != "10") & (p["procedure_modifier_2"] != "95")]
    if runout:
        p = p[p["paid"] <= RUNOUT]
    p = _versions(p, versions)
    if specialty == "provider":
        prov = X["provfile"]
        tax = prov.groupby("provider_npi")["primary_taxonomy"].agg(lambda x: sorted(set(x))[0]).to_dict()
        t = p["rendering_provider_npi"].map(tax)
    else:
        t = p["rendering_provider_taxonomy"]
    p = p.assign(tax=t.values)
    if physicians:
        p = p[p["tax"].str[:3].isin(PHYS_TAX)]
    else:
        p = p.assign(tax=p["tax"].where(p["tax"].str[:3].isin(PHYS_TAX), "207Q00000X"))
    p = p.assign(spec=p["tax"].map(SPECIALTY))
    if system == "register":
        ldg = [ctrl.system_of(tin, d) == "Ledgestone Health" for tin, d in zip(p["billing_provider_tin"], p["dos"])]
    else:
        lmg = ctrl_tin(X, "Ledgestone Medical Group")
        ldg = list(p["billing_provider_tin"] == lmg)
    p = p.assign(ldg=ldg, bmg=(p["billing_provider_tin"] == bmg_tin).values)
    key = ["member_id", "dos", "procedure_code", "rendering_provider_npi"]
    out = {}
    for sp in SPECIALTIES:
        sel = p[p["spec"] == sp]
        tot = _services(sel, key, second)
        b = _services(sel[sel["bmg"]], key, second)
        lg = _services(sel[sel["ldg"]], key, second)
        out[sp] = (tot, b, lg, 100.0 * b / tot, 100.0 * lg / tot)
    return out


def ctrl_tin(X, name):
    reg = X["register"]
    return reg.loc[reg["entity_name"] == name, "entity_tin"].iloc[0]


# ------------------------------------------------------------------ the ladder, from the shipped files
def ldg_2024_prices(X):
    """Ledgestone's 2024 allowed amount per unit by carrier and code, as its claims show it."""
    fac = latest_versions(X["fac"])
    facdir = X["facdir"]
    ldg = set(facdir.loc[facdir["health_system"] == "Ledgestone Health", "facility_npi"])
    car_name = dict(zip(X["prodref"]["submitter_code"], X["prodref"]["carrier_name"]))
    f = fac[(fac["dos"].dt.year == 2024) & fac["service_facility_npi"].isin(ldg) & (fac["qty"] > 0)]
    out = {}
    for c, cpt, a, q in zip(f["submitter_code"], f["procedure_code"], f["allowed"], f["qty"]):
        out.setdefault((car_name[c], cpt), round(a / q, 2))
    return out


def group_share(X, g, tin):
    """2024 share of a group's referred services (codes Ledgestone offers) that went to Ledgestone."""
    fac = latest_versions(X["fac"])
    spans = billing_spans(X["reassign"])
    facdir = X["facdir"]
    sys_of = dict(zip(facdir["facility_npi"], facdir["health_system"]))
    f = fac[(fac["dos"].dt.year == 2024) & (fac["paid"] <= RUNOUT) & fac["procedure_code"].isin(g["offered"])]
    sel = [any(t == tin and s <= d <= e for t, s, e in spans.get(n, ())) for n, d in
           zip(f["referring_provider_npi"], f["dos"])]
    f = f[sel]
    return float((f["service_facility_npi"].map(sys_of) == "Ledgestone Health").mean()), len(f)


def ladder(target, g):
    X = load(target)
    br = g["br"]
    el = br[br["eligible"]]
    P = g["P"]
    h = g["hist"]
    fit_all = g["fit"]
    ldg_tins = h.loc[h["acq_sys"] == "Ledgestone Health", "prac_tin"].unique()
    fit_ldg = fit_models(h[h["prac_tin"].isin(ldg_tins)])
    # every estimate is stated in whole per cent before it is projected (4.05)
    fit_all = {m: {k: whole(v) for k, v in d.items()} for m, d in fit_all.items()}
    fit_ldg = {m: {k: whole(v) for k, v in d.items()} for m, d in fit_ldg.items()}
    s_bmg = float(br.loc[br["offered"], "at_ldg"].mean())
    old = ldg_2024_prices(X)
    prices = {"filed": el["new_rate"].tolist(),
              "claims2024": [old.get((c, cpt), nr) for c, cpt, nr in zip(el["carrier"], el["procedure_code"], el["new_rate"])]}
    allbase = 1.0 / (1.0 - s_bmg)

    def proj(capture, price="filed", base=1.0):
        caps = capture(el)
        return P + base * sum(c * (pr * q - a) for c, pr, q, a in zip(caps, prices[price], el["qty"], el["allowed"]))

    prod_ldg = fit_ldg["product"]
    models = {
        "pooled, Ledgestone's purchases": lambda df: [fit_ldg["pooled"]["all"]] * len(df),
        "pooled, every acquisition": lambda df: [fit_all["pooled"]["all"]] * len(df),
        "by carrier": lambda df: [fit_ldg["carrier"].get(c, 0.0) for c in df["submitter_code"]],
        "product by product": lambda df: [prod_ldg.get(p, 0.0) for p in df["product"]],
        "tier, 2024 placements": lambda df: [fit_all["tier"]["1"] if t == 1 else fit_all["tier"]["2"] for t in df["t24"]],
        "tier, 2026 placements": lambda df: [fit_all["tier"]["1"] if t == 1 else fit_all["tier"]["2"] for t in df["t26"]],
    }
    grid = {}
    for name, cap in models.items():
        for price in ("filed", "claims2024"):
            for bname, b in (("outside", 1.0), ("all", allbase)):
                grid[(name, price, bname)] = proj(cap, price, b)
    lmg_share, n_lmg = group_share(X, g, tin_of_name(X, "Ledgestone Medical Group"))
    mfp_share, n_mfp = group_share(X, g, tin_of_name(X, "Mishicot Family Practice"))
    c_lmg = whole((lmg_share - s_bmg) / (1 - s_bmg))
    c_mfp = whole((mfp_share - s_bmg) / (1 - s_bmg))
    rungs = {
        "R0": P,
        "R1 employed group": proj(lambda df: [c_lmg] * len(df)),
        "R1 lakeshore affiliate": proj(lambda df: [c_mfp] * len(df)),
        "R2 pooled, Ledgestone": grid[("pooled, Ledgestone's purchases", "filed", "outside")],
        "R2 pooled, every acquisition": grid[("pooled, every acquisition", "filed", "outside")],
        "R3 carrier": grid[("by carrier", "filed", "outside")],
        "R4 product by product": grid[("product by product", "filed", "outside")],
        "R4 product by product, every acquisition": proj(lambda df: [fit_all["product"].get(p, 0.0) for p in df["product"]]),
        "R4 product and system": proj(lambda df: [fit_all["product x system"].get("%s|Ledgestone Health" % p, 0.0)
                                                  for p in df["product"]]),
        "R4 tier, 2024 placements": grid[("tier, 2024 placements", "filed", "outside")],
        "R5 answer": grid[("tier, 2026 placements", "filed", "outside")],
    }
    fam24 = {f: float(sum((fit_all["tier"]["1"] if t == 1 else fit_all["tier"]["2"]) * d
                         for t, d, ff in zip(el["t24"], el["delta"], el["fam"]) if ff == f)) for f in LINES}
    famr4 = {f: float(sum(prod_ldg.get(p, 0.0) * d for p, d, ff in zip(el["product"], el["delta"], el["fam"]) if ff == f))
             for f in LINES}
    return dict(rungs=rungs, grid=grid, s_bmg=s_bmg, c_lmg=c_lmg, c_mfp=c_mfp, n_lmg=n_lmg, n_mfp=n_mfp,
                fit_ldg=fit_ldg, fam24=fam24, famr4=famr4)


def tin_of_name(X, name):
    reg = X["register"]
    return reg.loc[reg["entity_name"] == name, "entity_tin"].iloc[0]
