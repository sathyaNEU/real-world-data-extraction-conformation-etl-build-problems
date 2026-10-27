"""Independent verifier for task258: reads only target/ and golden/, on its own parsing path (the csv module,
zoneinfo, openpyxl, pypdf), recomputes the call, every property's page loads moved, the default path's winner, both
asks, and checks every owed figure in the golden assignment paper.

    python3 task258/generator/verify.py [target] [golden]"""
from __future__ import annotations

import csv
import datetime as dt
import math
import os
import re
import statistics
import sys
from collections import defaultdict
from zoneinfo import ZoneInfo

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
SYD = ZoneInfo("Australia/Sydney")
UTC = dt.timezone.utc


def rows(target, name):
    with open(os.path.join(target, name), newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def find(target, pattern):
    hits = [f for f in os.listdir(target) if re.fullmatch(pattern, f)]
    assert len(hits) == 1, (pattern, hits)
    return hits[0]


def utc(s):
    return dt.datetime.strptime(s, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=UTC)


def half_up(x, unit):
    return int(math.floor(x / unit + 0.5) * unit)


def main(target=None, golden=None):
    target = target or os.path.join(TASK, "target")
    golden = golden or os.path.join(TASK, "golden")
    month0 = dt.datetime(2026, 11, 1, tzinfo=SYD).astimezone(UTC)
    month1 = dt.datetime(2026, 12, 1, tzinfo=SYD).astimezone(UTC)

    rates = defaultdict(list)
    for r in rows(target, find(target, r"rum_sampling_config\.csv")):
        rates[r["origin"]].append((utc(r["effective_from_utc"]), float(r["sample_rate"])))
    for o in rates:
        rates[o].sort()

    def rate(origin, t):
        out = None
        for since, v in rates[origin]:
            if since <= t:
                out = v
        return out

    hosts = rows(target, find(target, r"cdn_host_list\.csv"))
    prop_of = {h["host"]: h["property"] for h in hosts if h["kind"] == "page origin"}
    own = defaultdict(set)
    cdn = {}
    for h in hosts:
        if h["kind"] in ("image host", "legacy image alias"):
            own[h["property"]].add(h["host"])
        elif h["kind"] == "image CDN path":
            cdn[h["property"]] = h["path_prefix"]
    six = [p for p in dict.fromkeys(prop_of[o] for o in rates)]

    import openpyxl
    ws = openpyxl.load_workbook(os.path.join(target, find(target, r"squad_change_log_.*\.xlsx")),
                                data_only=True)["Engagements"]
    win = {}
    for r in ws.iter_rows(values_only=True):
        if isinstance(r[0], str) and re.fullmatch(r"E\d\d", r[0]):
            win[r[0]] = [v.date() for v in r[4:9]]
    parts = ("ttfb_ms", "load_delay_ms", "load_time_ms", "render_delay_ms")
    seen = {"before": defaultdict(list), "after": defaultdict(list)}
    by = defaultdict(lambda: defaultdict(list))
    for r in rows(target, find(target, r"squad_engagement_beacons_.*\.csv")):
        d = utc(r["fired_at_utc"]).astimezone(SYD).date()
        b0, b1, ship, a0, a1 = win[r["engagement"]]
        w = "before" if b0 <= d <= b1 else "after" if a0 <= d <= a1 else None
        if w:
            for c in parts:
                seen[w][c].append(float(r[c]))
            by[(r["engagement"], w)]["load_delay_ms"].append(float(r["load_delay_ms"]))
    level = statistics.median(seen["after"]["load_delay_ms"])
    r_load = statistics.median(seen["after"]["load_time_ms"]) / statistics.median(seen["before"]["load_time_ms"])
    r_delay = statistics.median(seen["after"]["load_delay_ms"]) / statistics.median(seen["before"]["load_delay_ms"])
    after_meds = [statistics.median(v["load_delay_ms"]) for k, v in by.items() if k[1] == "after"]

    # the nightly crawl: each page's <img> sources as its server sent them
    in_html = set()
    for r in rows(target, find(target, r"site_crawl_.*\.csv")):
        in_html.add((r["page_url"], r["img_src"]))

    moved, held_path, ratio_path, zero_path, field = (defaultdict(float) for _ in range(5))
    ids = set()
    for r in rows(target, find(target, r"rum_lcp_beacons_.*\.csv")):
        if r["beacon_id"] in ids:
            continue
        ids.add(r["beacon_id"])
        t = utc(r["fired_at_utc"])
        if not (month0 <= t < month1) or r["lcp_ms"] == "":
            continue
        name = prop_of[r["origin"]]
        url = r["lcp_url"]
        mine = r["lcp_element"] == "IMG" and url != "" and (
            url.split("/")[2] in own[name] or url.startswith("https://img.brolgacdn.au" + cdn[name]))
        if not mine or int(r["lcp_ms"]) <= 2500:
            continue
        w = 1.0 / rate(r["origin"], t)
        tt, dd, uu, qq = (float(r[c]) for c in parts)
        fcp = float(r["fcp_ms"])
        # a preload in the document head can only name a hero the page's server HTML carries; where it does not,
        # the image's request still starts when it starts today
        reach = ("https://" + r["origin"] + r["page_path"], url) in in_html
        delay = level if reach else dd
        if max(tt + delay + r_load * uu + qq, fcp) <= 2500:
            moved[name] += w
        if max(tt + level + r_load * uu + qq, fcp) <= 2500:
            field[name] += w
        if tt + level + r_load * uu + qq <= 2500:
            held_path[name] += w
        if tt + r_delay * dd + r_load * uu + qq <= 2500:
            ratio_path[name] += w
        if tt + r_load * uu + qq <= 2500:
            zero_path[name] += w
    gold = {n: half_up(moved[n], 1e5) for n in six}
    call = max(six, key=lambda n: moved[n])
    ranked = sorted(six, key=lambda n: -moved[n])
    print("verify: level %.1f ms (engagement after-delay medians %d to %d ms), load time ratio %.4f; loads moved %s"
          % (level, min(after_meds), max(after_meds), r_load, ", ".join("%s %.1fM" % (n, gold[n] / 1e6) for n in ranked)))
    print("verify: call %s (%.3fM, %.2fx the runner-up); ratios of medians file %s, a zero delay files %s, the level "
          "with the render delay held files %s, the level with the first-paint floor on every load files %s"
          % (call, moved[call] / 1e6, moved[ranked[0]] / moved[ranked[1]], max(six, key=lambda n: ratio_path[n]),
             max(six, key=lambda n: zero_path[n]), max(six, key=lambda n: held_path[n]),
             max(six, key=lambda n: field[n])))

    # Ask A
    import openpyxl as ox
    av = {}
    wa = ox.load_workbook(os.path.join(target, find(target, r"squad_avif_size_table\.xlsx")), data_only=True)["AVIF sizes"]
    for r in wa.iter_rows(min_row=5, values_only=True):
        if r[0] and r[1] == "hero":
            av[r[0]] = ("/i/%dx%d/" % (r[2], r[3]), float(r[4]))
    from pypdf import PdfReader
    text = " ".join(p.extract_text() for p in PdfReader(os.path.join(target, find(target, r"image_cdn_contract.*\.pdf")))
                    .pages)
    aud_gb = float(re.search(r"AUD (\d+\.\d+) per GB", text).group(1))
    req = defaultdict(int)
    for r in rows(target, find(target, r"image_host_logs_.*\.csv")):
        t = dt.datetime.strptime(r["hour_utc"], "%Y-%m-%d %H:%M").replace(tzinfo=UTC)
        if not (month0 <= t < month1) or r["status"] != "200":
            continue
        for n in six:
            hero_dir, _ = av[n]
            if (r["host"] in own[n] and r["path_prefix"] == hero_dir) or \
                    (r["host"] == "img.brolgacdn.au" and r["path_prefix"] == cdn[n].rstrip("/") + hero_dir):
                req[n] += int(r["requests"])
    for r in rows(target, find(target, r"legacy_media_server_log_.*\.csv")):
        t = dt.datetime.strptime(r["log_hour"], "%d/%b/%Y:%H %z").astimezone(UTC)
        if not (month0 <= t < month1) or r["code"] != "200":
            continue
        for n in six:
            if r["vhost"] in own[n] and r["uri_dir"] == av[n][0]:
                req[n] += int(r["hits"])
    gb = {n: req[n] * av[n][1] / 1e9 for n in six}
    aud = {n: gb[n] * aud_gb for n in six}

    # Ask B
    dam = rows(target, find(target, r"dam_image_export_.*\.csv"))
    latest = max(r["last_used_on"] for r in dam)
    live = [r for r in dam if r["last_used_on"] == latest]
    legacy = {r["legacy_id"] for r in live if r["legacy_id"]}
    assets = defaultdict(set)
    for r in live:
        if r["asset_id"] not in legacy:
            assets[prop_of[r["last_used_page"].split("/")[2]]].add(r["asset_id"])
    heroes = {n: len(assets[n]) for n in six}

    # the golden paper
    note = open(os.path.join(golden, "performance_squad_q1_2027.md"), encoding="utf-8").read()
    table = {}
    for line in note.splitlines():
        m = re.match(r"\| (\w+)(?: \(squad\))? \| ([\d.]+) million \| ([\d,]+) \| \$([\d,]+) \| ([\d,]+) \|", line)
        if m:
            table[m.group(1)] = (float(m.group(2)), int(m.group(3).replace(",", "")), int(m.group(4).replace(",", "")),
                                 int(m.group(5).replace(",", "")))
    bad = []
    for n in six:
        want = (round(gold[n] / 1e6, 1), half_up(gb[n], 1), half_up(aud[n], 1), heroes[n])
        if table.get(n) != want:
            bad.append((n, table.get(n), want))
    assert not bad, bad
    assert note.splitlines()[0].endswith(call) and ("**Send the squad to %s.**" % call) in note
    print("verify: egress %s" % ", ".join("%s %d GB $%d" % (n, half_up(gb[n], 1), half_up(aud[n], 1)) for n in six))
    print("verify: live hero images %s" % ", ".join("%s %d" % (n, heroes[n]) for n in six))
    print("verify: the golden paper agrees on the call and 24 figures")
    return call, gold, gb, aud, heroes


if __name__ == "__main__":
    main(*(sys.argv[1:3]))
