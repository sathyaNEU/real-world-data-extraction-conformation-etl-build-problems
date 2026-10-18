"""task243 goldens: the review workbook (each city and municipality's death registration completeness, its range over
the fitting choices, whether it goes into the first round of CRIP 2024, and the two indicators beside it) and the note
to the Provincial Civil Registration Committee with its chart."""
import io
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import postprocess

XLSX = "lcro_death_registration_completeness_2015_2020.xlsx"
DOCX = "crvs_committee_note_2023.docx"
ORG = "Kalinawan State University"
AUTHOR = "Nicole Gomez"


def _workbook(path, A):
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    ws["A1"] = "Death registration completeness, 1 August 2015 to 30 April 2020, and the CRIP 2024 first round"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Province of Kalinawan: review for the Provincial Civil Registration Committee, meeting of 21 November "
                "2023, by the Population Studies Center, Kalinawan State University. Indicators as defined in the terms "
                "of reference of 9 October 2023.")
    ws["A2"].font = Font(italic=True, size=9)
    head = ["City / municipality", "Registered deaths of residents aged 15 and over, 1 Aug 2015 to 30 Apr 2020",
            "Completeness of death registration, D1 (%)", "D1 over the fitting choices (%)",
            "CRIP 2024 first round (D1 below 80%)", "Births registered within 30 days, 2022, B1 (%)",
            "Deaths certified by a licensed physician, 2022, C1 (%)"]
    thin = Side(style="thin", color="A6ADB8")
    for j, h in enumerate(head, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(bold=True, size=9.5)
        c.fill = PatternFill("solid", fgColor="DDE4EE")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = Border(bottom=thin)
    ws.row_dimensions[4].height = 58
    for i, r in enumerate(A["lgus"], 5):
        vals = [r["name"], r["reg_adult"], r["completeness"], "%d to %d" % tuple(r["interval"]),
                "Goes in" if r["alert"] else "", r["b1"], r["c1"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(size=10, bold=(j == 5 and r["alert"]))
            if j == 2:
                c.number_format = "#,##0"
            elif j == 3:
                c.number_format = "0"
            elif j in (6, 7):
                c.number_format = "0.0"
            c.alignment = Alignment(horizontal="left" if j == 1 else ("center" if j in (4, 5) else "right"))
            if r["alert"]:
                c.fill = PatternFill("solid", fgColor="FBEFD9")
    n = 5 + len(A["lgus"])
    notes = [
        "D1: registered deaths of usual residents aged 15 and over, as a share of all their deaths in the interval, "
        "estimated with the general growth balance on the 2015 POPCEN and 2020 CPH household population, with each "
        "city's and municipality's net movement above each age from the 2020 CPH residence five years ago (fitted over "
        "ages 15+ to 65+). The range shows the same balance over the other usual fitting ranges and choices.",
        "Registered deaths: certificates of death in the CRS extract of 31 March 2023, a death registered at two offices "
        "counted once; deaths with age not stated (%d in the interval) are spread over the stated ages." %
        A["counts"]["window_unstated_age"],
        "B1 and C1 are for the committee's second round and do not change the first round.",
        "Sources: CRS extract (PSA provincial office, 31 March 2023); PSA, 2015 POPCEN and 2020 CPH; 2020 CPH special "
        "tabulation of residence five years ago (12 October 2023); PHO roster of health office clinical staff."]
    for k, t in enumerate(notes):
        c = ws.cell(row=n + 1 + k, column=1, value=t)
        c.font = Font(size=8.5, italic=(k == len(notes) - 1))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=n + 1 + k, start_column=1, end_row=n + 1 + k, end_column=7)
        ws.row_dimensions[n + 1 + k].height = 40 if k == 0 else 26
    for col, w in zip("ABCDEFG", [22, 18, 15, 14, 16, 15, 16]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B5"
    ws.print_title_rows = "4:4"

    gs = wb.create_sheet("Growth balance")
    gs["A1"] = "General growth balance by city and municipality (ages 15+ to 65+), with net movement"
    gs["A1"].font = Font(bold=True, size=11)
    gh = ["City / municipality", "Household population 15+, 2015", "Household population 15+, 2020",
          "Net movers aged 50+, 1 May 2015 to 1 May 2020", "Slope", "Intercept", "Completeness (%)"]
    for j, h in enumerate(gh, 1):
        c = gs.cell(row=3, column=j, value=h)
        c.font = Font(bold=True, size=9.5)
        c.fill = PatternFill("solid", fgColor="DDE4EE")
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    gs.row_dimensions[3].height = 44
    for i, r in enumerate(A["lgus"], 4):
        vals = [r["name"], r["pop15_2015"], r["pop15_2020"], r["movers50"], r["slope"], r["intercept"],
                round(100 * r["raw"], 1)]
        for j, v in enumerate(vals, 1):
            c = gs.cell(row=i, column=j, value=v)
            c.number_format = {2: "#,##0", 3: "#,##0", 4: "#,##0;-#,##0", 5: "0.000", 6: "0.0000", 7: "0.0"}.get(j, "General")
    for col, w in zip("ABCDEFG", [22, 16, 16, 18, 10, 11, 14]):
        gs.column_dimensions[col].width = w
    gs.freeze_panes = "B4"

    ns = wb.create_sheet("Notes")
    lines = [
        ("As of", "6 November 2023 (CRS extract of 31 March 2023)"),
        ("Interval", "1 August 2015 (2015 POPCEN reference date) to 30 April 2020 (the day before the 2020 CPH reference date)"),
        ("Deaths", "Certificates of death (fetal deaths excluded) of the province's usual residents, by the residence on the "
                   "certificate; a certificate repeated in the extract counted once, a death registered at two offices "
                   "counted once at its first registration."),
        ("Movement", "Persons in the city or municipality in 2020 who lived elsewhere on 1 May 2015, less persons anywhere "
                     "in the country in 2020 who lived there then; each age group's movers placed half a group younger "
                     "(their age when they moved), carried from five years to the 4.75-year interval."),
        ("Alert", "D1 below 80 per cent: the office goes into the first round of CRIP 2024."),
        ("B1", "Each child counted once, at its first registration; twins are two births."),
        ("C1", "A certificate made in the health office's name counts as certified by a physician when the person who "
               "signed it held a physician's post on the date of death (PHO roster)."),
    ]
    for i, (k, v) in enumerate(lines, 1):
        ns.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = ns.cell(row=i, column=2, value=v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ns.row_dimensions[i].height = 30 if len(v) > 100 else 16
    ns.column_dimensions["A"].width = 12
    ns.column_dimensions["B"].width = 110
    wb.save(path)
    postprocess.scrub(path, AUTHOR, "2023-11-14")
    postprocess.finish_office(path, "2023-11-07T00:40:00Z", "2023-11-14T09:12:00Z", ORG)


def chart_png(A):
    rows = sorted(A["lgus"], key=lambda r: r["completeness"])
    names = [r["name"] for r in rows]
    vals = [r["completeness"] for r in rows]
    lo = [r["completeness"] - r["interval"][0] for r in rows]
    hi = [r["interval"][1] - r["completeness"] for r in rows]
    cols = ["#B5532B" if r["alert"] else "#7A8CA5" for r in rows]
    fig, ax = plt.subplots(figsize=(6.6, 4.4), dpi=200)
    y = range(len(rows))
    ax.barh(list(y), vals, color=cols, height=0.62, xerr=[lo, hi], error_kw=dict(ecolor="#333333", lw=0.8, capsize=2.5))
    ax.axvline(80, color="#1F2A3A", lw=1.2, ls="--")
    ax.text(80.8, -0.95, "Alert: 80 per cent", fontsize=7.5, color="#1F2A3A", va="center")
    for i, r in enumerate(rows):
        ax.text(1.5, i, "%d%%" % r["completeness"], va="center", ha="left", fontsize=7.5, color="white", fontweight="bold")
        if r["alert"]:
            ax.text(r["interval"][1] + 1.2, i, "Round 1", va="center", fontsize=7, color="#B5532B")
    ax.set_yticks(list(y))
    ax.set_yticklabels(names, fontsize=8)
    ax.set_xlim(0, 106)
    ax.set_ylim(-1.5, len(rows) - 0.4)
    ax.set_xticks([0, 20, 40, 60, 80, 100])
    ax.set_xlabel("Completeness of death registration, residents aged 15 and over, Aug 2015 to Apr 2020 (%)", fontsize=7.8)
    ax.set_title("Five offices register under 80 per cent of their residents' adult deaths", fontsize=9.5, loc="left",
                 fontweight="bold")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.tick_params(axis="x", labelsize=7.5)
    fig.text(0.01, 0.012, "Bars: general growth balance with net movement, ages 15+ to 65+; whiskers: the other usual fitting "
                          "ranges and choices.\nSource: CRS extract of 31 March 2023; PSA, 2015 POPCEN, 2020 CPH and its "
                          "special tabulation of residence five years ago.", fontsize=5.8, color="#555555", va="bottom")
    fig.tight_layout(rect=(0, 0.06, 1, 1))
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def tens(v):
    return "{:,}".format(int(abs(v) / 10.0 + 0.5) * 10)


def span(*v):
    return ("%d" % v[0]) if min(v) == max(v) else ("%d to %d" % (min(v), max(v)))


def _note(path, A):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    hdr = sec.header.paragraphs[0]
    hdr.text = "Population Studies Center, Kalinawan State University"
    hdr.runs[0].font.size = Pt(8)
    ftr = sec.footer.paragraphs[0]
    ftr.text = "For the Provincial Civil Registration Committee, 21 November 2023"
    ftr.runs[0].font.size = Pt(8)

    def para(text, bold=False, size=None, after=6, italic=False, color=None):
        p = d.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        p.paragraph_format.space_after = Pt(after)
        return p

    des = [r for r in A["lgus"] if r["alert"]]
    names = [r["name"] for r in des]
    para("NOTE TO THE PROVINCIAL CIVIL REGISTRATION COMMITTEE", bold=True, size=9, after=2, color="555555")
    para("From: Nicole Gomez, Research Associate, Population Studies Center    Date: 14 November 2023    Re: CRIP 2024 "
         "first round", size=9,
         after=10, color="555555")
    the = lambda n: ("the " + n) if n.startswith("City of") else n
    lead = ", ".join(the(n) for n in names[:-2]) + ", %s and %s" % (the(names[-2]), the(names[-1]))
    para("%s register fewer than 80 per cent of their residents' adult deaths and go into the first round" %
         (lead[0].upper() + lead[1:]), bold=True, size=13, after=8)
    rd = {r["name"]: r for r in A["lgus"]}
    para("Between the 2015 POPCEN and the 2020 CPH the five offices registered the deaths of about %d to %d per cent of "
         "their residents who died at 15 or over (indicator D1 in the terms of reference); every other office registered "
         "%d per cent or more. The chart shows each city and municipality against the alert, with the range the estimate "
         "takes over the other usual ways of fitting it." % (min(r["completeness"] for r in des), max(r["completeness"] for r in des),
                                                               min(r["completeness"] for r in A["lgus"] if not r["alert"])),
         after=8)
    d.add_picture(io.BytesIO(chart_png(A)), width=Inches(6.4))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    ep, mg, sf, ct = rd["Esperanza"], rd["Magsaysay"], rd["Mabini"], rd["Kalinawan City"]
    para("How completeness was read. There is no complete count of deaths to set the registered ones against, so D1 is "
         "estimated from the two censuses: between them the population above each age changes by the people who reach "
         "that age, the deaths above it and the people who move in or out above it, and the deaths the registers should "
         "hold follow from that balance (the general growth balance). Kalinawan's older residents move in numbers: in "
         "the five years to the 2020 census Esperanza lost a net of about %s people aged 50 and over, Magsaysay about "
         "%s and Mabini about %s, most of them to other parts of the country, while Kalinawan City gained about %s. "
         "Read as if nobody moved, the balance takes each departure for a death nobody registered and would put all "
         "three in the seventies, under the alert. With the movement from the 2020 census's question on residence five "
         "years ago in the balance, the residents who left the province included, the three farming municipalities "
         "register %s per cent." %
         (tens(-ep["movers50"]), tens(-mg["movers50"]), tens(-sf["movers50"]), tens(ct["movers50"]),
          span(ep["completeness"], mg["completeness"], sf["completeness"])),
         after=8)
    b1 = sorted(A["lgus"], key=lambda r: r["b1"])
    c1 = sorted(A["lgus"], key=lambda r: r["c1"])
    para("For the second round. Timely birth registration (B1) in 2022 ran from %.1f per cent in %s to %.1f per cent in "
         "%s, counting each child once at its first registration. Medical certification of 2022 deaths (C1) ran from "
         "%.1f per cent in %s to %.1f per cent in %s: a certificate made in the health office's name counts only where "
         "the person who signed it held a physician's post that day, and in four offices the health officer's post was "
         "vacant for part of 2022. The workbook (%s) has every figure." %
         (b1[0]["b1"], the(b1[0]["name"]), b1[-1]["b1"], the(b1[-1]["name"]), c1[0]["c1"], the(c1[0]["name"]),
          c1[-1]["c1"], the(c1[-1]["name"]), XLSX), after=8)
    para("Sources: CRS extract run by the PSA provincial office as of 31 March 2023; PSA, 2015 POPCEN and 2020 CPH household "
         "population; 2020 CPH special tabulation of residence five years ago (12 October 2023); PHO roster of the health "
         "offices' clinical staff (15 July 2023).", italic=True, size=8.5, after=0)
    d.core_properties.author = AUTHOR
    d.core_properties.title = "Note to the Provincial Civil Registration Committee: CRIP 2024 first round"
    d.core_properties.comments = ""
    d.save(path)
    postprocess.scrub(path, AUTHOR, "2023-11-14")
    postprocess.finish_office(path, "2023-11-09T01:05:00Z", "2023-11-14T09:30:00Z", ORG)


def write(folder, A):
    os.makedirs(folder, exist_ok=True)
    _workbook(os.path.join(folder, XLSX), A)
    _note(os.path.join(folder, DOCX), A)
