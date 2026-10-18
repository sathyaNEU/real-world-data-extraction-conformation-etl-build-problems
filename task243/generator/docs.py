"""task243 documents: the provincial administrator's request (eml), the committee's terms of reference (pdf), the
census office's cover note for the special tabulation (pdf), the two censuses' household population tables (xlsx),
the provincial health office's roster of health officers (xlsx), the LCRO capacity survey (xlsx), the data dictionary
and the readme (md)."""
import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess
import records as RC

PROVGOV = "Provincial Government of Kalinawan"
PEOPLE = {"admin": ("Steven Jimenez", "steven.jimenez", "kalinawan.gov.ph"),
          "popn": ("Nicole Gomez", "nicole.gomez", "kalinawan-su.edu.ph"),
          "registrar": ("Randy Cortez", "randy.cortez", "kalinawan.gov.ph")}


def _addr(k):
    name, user, domain = PEOPLE[k]
    return "%s <%s@%s>" % (name, user, domain)


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.6, leading=13.4, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=13.5, leading=17, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=14,
                        spaceBefore=8, spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10.5, textColor=colors.HexColor("#444444"))
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=8.4, leading=10.2)
    return body, h1, h2, small, cell


def _grid(rows, widths, cell):
    t = Table([[Paragraph(str(c), cell) for c in r] for r in rows], colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA0AB")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E3E8EF"))]))
    return t


def _footer(text):
    def foot(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(20 * mm, 10 * mm, text)
        canvas.drawRightString(A4[0] - 20 * mm, 10 * mm, "Page %d" % d.page)
        canvas.restoreState()
    return foot


# ------------------------------------------------------------------ the request (eml)
def request(path, F):
    body = f"""From: {_addr("admin")}
To: {_addr("popn")}
Cc: {_addr("registrar")}
Date: Mon, 30 Oct 2023 08:41:17 +0800
Subject: CRIP 2024 first round: death registration review for the 21 November committee
MIME-Version: 1.0
Content-Type: text/plain; charset="utf-8"
Content-Transfer-Encoding: 8bit

Nicole,

The Provincial Civil Registration Committee meets on Tuesday 21 November to name the offices that go into the first
round of the Civil Registration Improvement Programme in January. As agreed at the September meeting, the first round
goes to the local civil registry offices whose death registration falls below the completeness alert in the terms of
reference (attached as filed on 9 October), and the committee has asked the Center to carry out the review under the
Province's agreement with the University.

Please prepare the review for the committee: the workbook ({F["deliv_xlsx"]}) with each city and municipality, and a
short committee note ({F["deliv_docx"]}) with one chart the members can read at a glance. The committee also asked to
see the two other indicators in the terms of reference beside completeness, for the second round.

What you have (the PPDO has put the folder together):
- the Civil Registry System extract the PSA provincial office ran for the committee as of 31 March 2023 (deaths and
  births);
- the 2015 POPCEN and 2020 CPH household population tables for the province;
- the 2020 CPH special tabulation Casey's office sent the PPDO on 12 October for the PDPFP update, with its cover note;
- the PHO's roster of the city and municipal health offices' clinical staff;
- the city and municipality directory.
The FHSIS mortality reports and the July LCRO capacity survey are in the folder too, from the September meeting.

Randy will take the technical working group through the note on the 16th, so a draft by the 14th would help.

Thank you.

Steven Jimenez
Provincial Administrator
{PROVGOV}
"""
    with open(path, "w") as fh:
        fh.write(body)


# ------------------------------------------------------------------ the terms of reference (pdf)
def tor(path, F):
    body, h1, h2, small, cell = _styles()
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                            bottomMargin=18 * mm, title="Terms of reference: death registration review, CRIP 2024",
                            author="Provincial Civil Registration Committee", creator=PROVGOV,
                            subject="Review of the local civil registry offices for the Civil Registration Improvement Programme",
                            invariant=1)
    s = []
    s.append(Paragraph("Provincial Civil Registration Committee, Province of Kalinawan", small))
    s.append(Paragraph("Terms of reference: review of death registration for the 2024 Civil Registration Improvement "
                       "Programme", h1))
    s.append(Paragraph("Filed 9 October 2023 (Committee Resolution No. 2023-07)", small))
    s.append(Spacer(1, 6))
    s.append(Paragraph("1. Purpose", h2))
    s.append(Paragraph(
        "The Civil Registration Improvement Programme (CRIP) gives a local civil registry office (LCRO) a package of "
        "support for twelve months: a registration officer seconded from the province, barangay registration points, "
        "community reporting of deaths through the barangay health workers, and a supply of forms and equipment. The "
        "first round, from January 2024, goes to the offices whose registration of deaths falls below the completeness "
        "alert in section 3. This review measures each of the province's twelve cities and municipalities against it.",
        body))
    s.append(Paragraph("2. Scope", h2))
    s.append(Paragraph(
        "The LCROs of Kalinawan City, the City of Santa Cruz and the municipalities of San Isidro, Esperanza, Magsaysay, "
        "Mabini, Libertad, San Luis, Del Carmen, Quezon, Rizal and San Miguel. Each office is reviewed on the "
        "registration of its own residents' events, wherever in the country the event was registered.", body))
    s.append(Paragraph("3. Indicators", h2))
    rows = [["Code", "Indicator", "Definition"],
            ["D1", "Completeness of death registration",
             "The deaths of the city's or municipality's usual residents aged 15 years and over that occurred from 1 "
             "August 2015 (the reference date of the 2015 Census of Population) to 30 April 2020 (the day before the "
             "reference date of the 2020 Census of Population and Housing) and were registered in the Civil Registry "
             "System by 31 March 2023, as a share of all deaths of those residents in that interval. A death belongs to "
             "the city or municipality of the deceased's usual residence at death."],
            ["", "Completeness alert", "D1 below 80 per cent. The office goes into the first round of CRIP 2024."],
            ["B1", "Timely registration of births",
             "Live births in 2022 to mothers usually resident in the city or municipality that were registered within "
             "30 days of the birth, as a share of all live births in 2022 to those mothers registered in the Civil "
             "Registry System by 31 March 2023."],
            ["C1", "Medically certified deaths",
             "Deaths in 2022 of the city's or municipality's usual residents, registered in the Civil Registry System "
             "by 31 March 2023, whose cause of death was certified by a licensed physician, as a share of all those "
             "registered deaths."]]
    s.append(_grid(rows, [12 * mm, 42 * mm, 116 * mm], cell))
    s.append(Spacer(1, 4))
    s.append(Paragraph("B1 and C1 are reported beside D1 for the committee's second round and do not change the first "
                       "round's selection.", body))
    s.append(Paragraph("4. Sources", h2))
    s.append(Paragraph(
        "The Civil Registry System extract run by the PSA provincial office as of 31 March 2023 (certificates of death, "
        "fetal death and live birth for the province's residents and for every event registered at an office in the "
        "province); the household population tables of the 2015 Census of Population (POPCEN) and the 2020 Census of "
        "Population and Housing (CPH) for the province; any other census output the review needs.", body))
    s.append(Paragraph("5. Reporting", h2))
    s.append(Paragraph(
        "A workbook with one row per city and municipality and a short note to the committee with one chart. Figures "
        "the review estimates are shown with the precision the committee needs to make the selection.", body))
    doc.build(s, onFirstPage=_footer("Provincial Civil Registration Committee, Province of Kalinawan"),
              onLaterPages=_footer("Provincial Civil Registration Committee, Province of Kalinawan"))
    postprocess.scrub(path, "PCRC Secretariat, Kalinawan", "2023-10-09")


# ------------------------------------------------------------------ the census office's cover note (pdf)
def cover_note(path, F):
    body, h1, h2, small, cell = _styles()
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                            bottomMargin=18 * mm, title="2020 CPH special tabulation: residence five years ago",
                            author="PSA Provincial Statistical Office, Kalinawan", creator="PSA Kalinawan",
                            subject="Transmittal of a special tabulation to the Provincial Planning and Development Office",
                            invariant=1)
    s = []
    s.append(Paragraph("Philippine Statistics Authority, Provincial Statistical Office, Kalinawan", small))
    s.append(Paragraph("12 October 2023", small))
    s.append(Spacer(1, 6))
    s.append(Paragraph("2020 CPH special tabulation: residence five years ago, Province of Kalinawan", h1))
    s.append(Paragraph(
        "To the Provincial Planning and Development Coordinator. As requested for the update of the Provincial "
        "Development and Physical Framework Plan, we transmit the tabulation of the 2020 Census of Population and "
        "Housing (CPH) on residence five years ago for the province, file %s." % F["restab"], body))
    s.append(Paragraph("Contents", h2))
    s.append(Paragraph(
        "Persons five years old and over by city or municipality of usual residence on 1 May 2020 (the CPH reference "
        "date), city or municipality of usual residence on 1 May 2015, five-year age group on 1 May 2020, and sex. "
        "Residence five years ago was asked of every person five years old and over in the CPH basic questionnaire; "
        "the counts are from the full enumeration.", body))
    rows = [["Field", "Values"],
            ["CITY_MUN_2020", "Each city or municipality of the province, and ELSEWHERE IN THE PHILIPPINES for persons "
                              "enumerated outside the province whose residence on 1 May 2015 was in the province."],
            ["RESIDENCE_1MAY2015", "SAME CITY/MUNICIPALITY; the name of another city or municipality of the province; "
                                   "OTHER PROVINCE; FOREIGN COUNTRY; NOT STATED."],
            ["AGE_GROUP", "Age in completed years on 1 May 2020, five-year groups from 5 - 9 to 80 and over."],
            ["SEX", "MALE, FEMALE."], ["PERSONS", "Number of persons."]]
    s.append(_grid(rows, [40 * mm, 130 * mm], cell))
    s.append(Spacer(1, 6))
    s.append(Paragraph(
        "Persons in institutional households are not included. Rows with no persons are not listed. The rows for "
        "persons living elsewhere in the Philippines are compiled from the national enumeration and show only the "
        "city or municipality of the province where those persons lived in 2015.", body))
    s.append(Paragraph("Casey Carrillo, Chief Statistical Specialist", body))
    doc.build(s, onFirstPage=_footer("PSA Provincial Statistical Office, Kalinawan"),
              onLaterPages=_footer("PSA Provincial Statistical Office, Kalinawan"))
    postprocess.scrub(path, "PSA Kalinawan", "2023-10-12")


# ------------------------------------------------------------------ the census tables (xlsx)
def census(path, year, R):
    wb = Workbook()
    ws = wb.active
    ws.title = "Table 3"
    if year == "2015":
        ws["A1"] = ("Table 3. Household Population by Five-Year Age Group and Sex, by City/Municipality: "
                    "Province of Kalinawan, 2015")
        ws["A2"] = "2015 Census of Population (POPCEN), reference date 1 August 2015"
    else:
        ws["A1"] = ("Table 3. Household Population by Five-Year Age Group and Sex, by City/Municipality: "
                    "Province of Kalinawan, 2020")
        ws["A2"] = "2020 Census of Population and Housing (CPH), reference date 1 May 2020"
    ws["A1"].font = Font(bold=True, size=11)
    head = ["City/Municipality and Age Group", "Both Sexes", "Male", "Female"]
    for j, h in enumerate(head, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCE3EC")
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    row = 5
    for code, male, fem in R.census_table(year):
        name = RC.UPPER[code]
        if year == "2015" and code == "KLW-02":
            name = "SANTA CRUZ"
        ws.cell(row=row, column=1, value=name).font = Font(bold=True)
        row += 1
        ws.cell(row=row, column=1, value="All Ages")
        ws.cell(row=row, column=2, value=int(male.sum() + fem.sum()))
        ws.cell(row=row, column=3, value=int(male.sum()))
        ws.cell(row=row, column=4, value=int(fem.sum()))
        row += 1
        for g, lab in enumerate(RC.AGE_GROUPS):
            ws.cell(row=row, column=1, value="  " + lab)
            ws.cell(row=row, column=2, value=int(male[g] + fem[g]))
            ws.cell(row=row, column=3, value=int(male[g]))
            ws.cell(row=row, column=4, value=int(fem[g]))
            row += 1
        row += 1
    ws.cell(row=row, column=1, value="Note: Household population excludes persons in institutional households.")
    ws.cell(row=row + 1, column=1, value="Source: Philippine Statistics Authority, %s." %
            ("2015 Census of Population" if year == "2015" else "2020 Census of Population and Housing"))
    for r_ in range(5, row):
        for c_ in range(2, 5):
            ws.cell(row=r_, column=c_).number_format = "#,##0"
    ws.column_dimensions["A"].width = 38
    for col in "BCD":
        ws.column_dimensions[col].width = 13
    ws.freeze_panes = "A5"
    wb.save(path)
    stamp = "2016-06-30" if year == "2015" else "2022-07-14"
    postprocess.scrub(path, "Philippine Statistics Authority", stamp)
    postprocess.finish_office(path, stamp + "T09:00:00Z", stamp + "T09:00:00Z", "Philippine Statistics Authority")


def roster(path, R):
    wb = Workbook()
    ws = wb.active
    ws.title = "Health officers"
    ws["A1"] = "Provincial Health Office: city and municipal health offices, physicians, nurses and midwives, 2019 to 2023"
    ws["A1"].font = Font(bold=True, size=11)
    ws["A2"] = "As of 15 July 2023. One row per person and term in the post."
    df = R.roster()
    head = list(df.columns)
    for j, h in enumerate(head, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCE3EC")
    for i, rec in enumerate(df.itertuples(index=False), 5):
        for j, v in enumerate(rec, 1):
            ws.cell(row=i, column=j, value=v)
    for col, w in zip("ABCDEFG", [22, 22, 26, 34, 12, 12, 16]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    wb.save(path)
    postprocess.scrub(path, "Provincial Health Office", "2023-07-15")
    postprocess.finish_office(path, "2021-02-03T01:15:00Z", "2023-07-15T06:40:00Z", "Provincial Health Office, Kalinawan")


def capacity(path, R):
    wb = Workbook()
    ws = wb.active
    ws.title = "LCRO survey July 2023"
    ws["A1"] = "LCRO capacity survey, July 2023 (Provincial Civil Registration Committee secretariat)"
    ws["A1"].font = Font(bold=True, size=11)
    df = R.staffing()
    for j, h in enumerate(df.columns, 1):
        c = ws.cell(row=3, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCE3EC")
    for i, rec in enumerate(df.itertuples(index=False), 4):
        for j, v in enumerate(rec, 1):
            ws.cell(row=i, column=j, value=v)
    for col, w in zip("ABCDEFG", [30, 16, 16, 11, 14, 28, 28]):
        ws.column_dimensions[col].width = w
    wb.save(path)
    postprocess.scrub(path, "PCRC Secretariat", "2023-07-28")
    postprocess.finish_office(path, "2023-07-10T02:00:00Z", "2023-07-28T08:22:00Z", PROVGOV)


# ------------------------------------------------------------------ the dictionary and the readme (md)
def dictionary(F):
    return f"""# Data dictionary

Province of Kalinawan, death registration review for CRIP 2024. Dates are written YYYY-MM-DD.

## {F["deaths"]}

Civil Registry System (CRS) extract run by the PSA provincial office on 31 March 2023: one row per certificate held
for the deaths and fetal deaths of the province's residents, wherever in the country they were registered, and for
every death and fetal death registered at an office in the province. Certificates registered on or after 1 January 2015.

| Field | Meaning |
|---|---|
| REGISTRY_NO | The registering office's registry number (year of registration and sequence). |
| LCRO_CODE, LCRO_NAME | The registering local civil registry office (PSGC code and name). |
| DATE_REGISTERED | Date the office registered the certificate. |
| CERT_TYPE | DEATH (Certificate of Death, Municipal Form No. 103) or FETAL DEATH (Certificate of Fetal Death, Municipal Form No. 103A). |
| LAST_NAME, FIRST_NAME, MIDDLE_NAME | The deceased's name as registered. |
| SEX | MALE or FEMALE. |
| DATE_OF_DEATH | Date of death, or of delivery for a fetal death. |
| DATE_OF_BIRTH | As written on the certificate; blank where the informant did not know it. |
| AGE, AGE_UNIT | Age at death as written on the certificate, in YEARS (completed years), MONTHS, DAYS or HOURS; blank where not stated. |
| CIVIL_STATUS | As registered. |
| RES_BARANGAY, RES_CITY_MUN, RES_PROVINCE | The deceased's usual residence as written on the certificate. |
| PLACE_OF_DEATH | HOSPITAL, HOME or OTHERS. |
| PLACE_OF_DEATH_CITY_MUN | City or municipality where the death occurred. |
| ATTENDANT | 1 private physician, 2 public health officer, 3 hospital authority, 4 none, 5 others. |
| CERTIFIED_BY, CERTIFIER_TITLE | The name and the title written under the signature in the certification of death. |
| LATE_REGISTRATION | Y where the certificate was registered more than 30 days after the death. |

## {F["births"]}

CRS extract of 31 March 2023: one row per certificate of live birth registered from 1 January 2022 to 31 March 2023 for
a birth to a mother usually resident in the province, wherever registered, or at an office in the province.

| Field | Meaning |
|---|---|
| REGISTRY_NO, LCRO_CODE, LCRO_NAME, DATE_REGISTERED | As in the death extract. |
| CHILD_FIRST_NAME, CHILD_MIDDLE_NAME, CHILD_LAST_NAME, SEX, DATE_OF_BIRTH | The child, as registered. |
| TYPE_OF_BIRTH, BIRTH_ORDER | SINGLE or TWIN; for a twin, FIRST or SECOND. |
| PLACE_OF_BIRTH, PLACE_OF_BIRTH_CITY_MUN | HOSPITAL, LYING-IN / BIRTHING HOME or HOME, and where. |
| MOTHER_FIRST_NAME, MOTHER_MAIDEN_LAST_NAME, MOTHER_AGE | The mother, age at the birth. |
| MOTHER_RES_BARANGAY, MOTHER_RES_CITY_MUN, MOTHER_RES_PROVINCE | The mother's usual residence as written. |
| FATHER_FIRST_NAME, FATHER_LAST_NAME | The father, where acknowledged. |
| ATTENDANT | PHYSICIAN, NURSE, MIDWIFE, HILOT or OTHERS. |
| LATE_REGISTRATION | Y where registered more than 30 days after the birth. |

## {F["cen2015"]}, {F["cen2020"]}

PSA tables of the household population by five-year age group and sex for each city and municipality: the 2015 Census
of Population (reference date 1 August 2015) and the 2020 Census of Population and Housing (reference date 1 May
2020). Each city or municipality has an All Ages row above its age groups. Santa Cruz became a component city on 16 March
2019 and is listed under its name of the day.

## {F["restab"]}

2020 CPH special tabulation of residence five years ago (see the cover note, {F["cover"]}): persons five years old and
over by residence on 1 May 2020 and on 1 May 2015, five-year age group on 1 May 2020 and sex.

| Field | Meaning |
|---|---|
| CITY_MUN_2020 | Residence on 1 May 2020: a city or municipality of the province, or ELSEWHERE IN THE PHILIPPINES. |
| RESIDENCE_1MAY2015 | SAME CITY/MUNICIPALITY, another city or municipality of the province, OTHER PROVINCE, FOREIGN COUNTRY or NOT STATED. |
| AGE_GROUP, SEX, PERSONS | Age group on 1 May 2020, sex, number of persons. |

## {F["directory"]}

The province's cities and municipalities: PSGC code, class, income class, number of barangays, the LCRO, and the date a
city's charter took effect.

## {F["roster"]}

Provincial Health Office roster of the clinical staff of each city and municipal health office from 2019 to July 2023: the
person, the position and the term (FROM, TO; TO blank for a person in post).

## {F["fhsis"]}

Field Health Services Information System (FHSIS) monthly reports, form M2: deaths reported to each city or municipal
health office by health facilities and barangay health workers, all ages and under one year, 2015 to 2022.

## {F["capacity"]}

The committee secretariat's July 2023 survey of the LCROs: staff, equipment, connectivity, documents awaiting encoding
and barangay registration points.
"""


def readme(F):
    return f"""# Province of Kalinawan: death registration review, CRIP 2024

Files gathered by the Provincial Planning and Development Office for the Provincial Civil Registration Committee's
meeting of 21 November 2023.

- {F["request"]}: the provincial administrator's request.
- {F["tor"]}: the committee's terms of reference for the review (indicators D1, B1, C1).
- {F["deaths"]}: Civil Registry System extract, deaths and fetal deaths, as of 31 March 2023.
- {F["births"]}: Civil Registry System extract, live births registered from January 2022, as of 31 March 2023.
- {F["cen2015"]}: 2015 POPCEN household population by age group and sex.
- {F["cen2020"]}: 2020 CPH household population by age group and sex.
- {F["restab"]} and {F["cover"]}: 2020 CPH special tabulation of residence five years ago, with the PSA cover note.
- {F["directory"]}: cities and municipalities and their LCROs.
- {F["roster"]}: Provincial Health Office roster of the health offices' clinical staff.
- {F["fhsis"]}: FHSIS monthly mortality reports.
- {F["capacity"]}: LCRO capacity survey, July 2023.
- DATA_DICTIONARY.md: the fields of each file.
"""


def write_all(target, F, R):
    import os
    p = lambda k: os.path.join(target, F[k])
    request(p("request"), F)
    tor(p("tor"), F)
    cover_note(p("cover"), F)
    census(p("cen2015"), "2015", R)
    census(p("cen2020"), "2020", R)
    roster(p("roster"), R)
    capacity(p("capacity"), R)
    with open(p("dict"), "w") as fh:
        fh.write(dictionary(F))
    with open(p("readme"), "w") as fh:
        fh.write(readme(F))
