"""task159 pack: every file the station planning desk pulled on Monday 11 January 2027, written from the records."""
from __future__ import annotations

import datetime as dt
import os

import numpy as np
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import common as K
import docs as DOC
import postprocess as PP
import records as R
import world as W

F_FARE = "vtc_fare_entries_hourly_2024-06-03_2027-01-10.parquet"
F_LEGACY = "afc_register_reads_2022-02-01_2024-06-30.csv"
F_DASH = "ridership_dashboard_avg_weekday_2025-2026.xlsx"
F_COMPLEX = "complex_reference.csv"
F_CENSUS = "station_census_2026-10.xlsx"
F_COUNTS = "platform_counts_autumn_2026.xlsx"
F_CAP = "platform_capacity_register.csv"
F_TT = "weekday_timetables_aut26_spr27.csv"
F_FLEET = "fleet_register.csv"
F_CHANGES = "service_changes_register.csv"
F_INV = "station_elements_inventory.csv"
F_ASSET = "escalator_asset_register.csv"
F_CTRL = "escalator_controller_log_2026-09-01_2026-12-31.csv"
F_WO = "elevator_work_orders_2026-09-01_2026-12-31.csv"
F_MON = "elevator_monitoring_2026-09-01_2026-12-31.csv"
F_DICT = "vtc_data_dictionary.xlsx"


def _ts(x):
    return x.strftime("%Y-%m-%d %H:%M:%S")


def _round_s(x):
    return (x + dt.timedelta(microseconds=500000)).replace(microsecond=0)


# ---------------------------------------------------------------------------------------------------- spreadsheets
HEAD_FILL = PatternFill("solid", fgColor="DCE6E1")
BOLD = Font(name="Calibri", bold=True, size=11)
BODY = Font(name="Calibri", size=11)


def _sheet(ws, head, rows, widths=None, fmt=None):
    ws.append(head)
    for c in ws[1]:
        c.font = BOLD
        c.fill = HEAD_FILL
    for r in rows:
        ws.append(list(r))
    for i, h in enumerate(head):
        ws.column_dimensions[get_column_letter(i + 1)].width = (widths[i] if widths else max(10, len(h) + 2))
    ws.freeze_panes = "A2"
    if fmt:
        for col, nf in fmt.items():
            for cell in ws[col][1:]:
                cell.number_format = nf


def write_census(out, L):
    path = os.path.join(out, F_CENSUS)
    wb = Workbook()
    ws = wb.active
    ws.title = "Counts"
    rows = R.census_rows(L)
    rows = [(r[0], r[1], dt.date.fromisoformat(r[2]), r[3], r[4], r[5], r[6], r[7]) for r in rows]
    _sheet(ws, ["Station", "Complex ID", "Census date", "Location ID", "Location", "Direction", "Quarter-hour",
                "People"], rows, [18, 11, 12, 12, 62, 14, 13, 9], fmt={"C": "yyyy-mm-dd"})
    ws2 = wb.create_sheet("Notes")
    notes = [
        ["Station census, October 2026"],
        ["Count programme, Station Planning. Counts by Tollison Field Data; supervised by %s." % K.PEOPLE["counts"]],
        [""],
        ["Each station was counted on one Tuesday, Wednesday or Thursday between 6 and 29 October 2026, 06:00 to 10:00."],
        ["People are counted in each quarter-hour at each platform access (the stairs and escalators at that end of the platform together) and in each passage between platforms."],
        ["Quarter-hours are labelled by their start, local time."],
        ["To platform: people going down to the platform from the mezzanine. From platform: people leaving the platform for the mezzanine. Passage: people walking from one platform to the other."],
        ["Island platforms are counted as one platform."],
    ]
    for n in notes:
        ws2.append(n)
    ws2.column_dimensions["A"].width = 120
    wb.save(path)
    PP.normalize_ooxml(path, "Caleb Rowe", "Caleb Rowe", "2026-11-23T14:05:00Z", "2026-11-30T21:41:00Z", excel=True,
                       company=K.ORG)
    return len(rows)


def count_label(key):
    station, rest = key.split("|")
    if len(rest) == 1:
        return "Line %s island" % rest
    return "Line %s %s" % (rest[0], W.DIR_WORD[rest[1:]])


def write_counts(out, L):
    path = os.path.join(out, F_COUNTS)
    rows = R.count_sheet_rows(L)
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Method"
    method = [
        ["Platform counts, autumn 2026"],
        ["Tollison Field Data for VTC Station Planning. Delivered 4 December 2026."],
        [""],
        ["At each train's arrival the enumerator records the number of people on the platform before the doors open. On an island platform the count is everyone on the island, and the train's direction is noted."],
        ["Crowds of about 200 or more are estimated to the nearest 10 and marked E in the Est column."],
        ["Due is the train's timetable arrival; Arrived is its arrival to the minute. All times are local."],
        ["Counts run from 06:45 to 09:50."],
    ]
    for m in method:
        ws0.append(m)
    ws0.column_dimensions["A"].width = 120
    for cx in W.WATCH_COMPLEXES:
        ws = wb.create_sheet(cx)
        sub = [r for r in rows if r[0] == cx]
        sub.sort(key=lambda r: (r[1], r[2], r[5], r[4]))
        out_rows = []
        for (st, d, key, dirn, due, arr, ppl, est) in sub:
            out_rows.append((d, count_label(key), dirn, due, arr, ppl, est))
        _sheet(ws, ["Date", "Platform", "Train", "Due", "Arrived", "People", "Est"], out_rows,
               [12, 30, 12, 8, 9, 9, 6], fmt={"A": "yyyy-mm-dd"})
    wb.save(path)
    PP.normalize_ooxml(path, "Tollison Field Data", "Caleb Rowe", "2026-10-13T16:20:00Z", "2026-12-04T19:02:00Z",
                       excel=True, company="Tollison Field Data")
    return len(rows)


def write_dashboard(out, fare):
    path = os.path.join(out, F_DASH)
    rows = R.dashboard_rows(fare)
    wb = Workbook()
    ws = wb.active
    ws.title = "Average weekday entries"
    _sheet(ws, ["Month", "Complex ID", "Complex", "Average weekday entries", "Status"], rows, [10, 11, 18, 22, 12],
           fmt={"D": "#,##0"})
    ws2 = wb.create_sheet("About")
    for n in [["Ridership dashboard: average weekday entries by complex"],
              ["Ridership Data, %s. Data owner: %s." % (K.ORG, K.PEOPLE["ridership"])],
              ["Entries are fare validations at the complex's arrays, every payment type, from the fare back office."],
              ["A service day runs from 03:00 to 02:59. Weekdays are Monday to Friday, federal holidays out. Averages are rounded to the nearest 10."],
              ["The latest month is preliminary until the 15th of the month after; earlier months are final."],
              ["Published 8 January 2027."]]:
        ws2.append(n)
    ws2.column_dimensions["A"].width = 110
    wb.save(path)
    PP.normalize_ooxml(path, K.PEOPLE["ridership"], K.PEOPLE["ridership"], "2025-02-10T15:00:00Z",
                       "2027-01-08T14:32:00Z", excel=True, company=K.ORG)
    return len(rows)


DICT_FILES = [
    (F_FARE, "Ridership Data (%s)" % K.PEOPLE["ridership"], "fare back office, hourly export", "3 June 2024 to 10 January 2027"),
    (F_LEGACY, "Ridership Data (%s)" % K.PEOPLE["ridership"], "legacy fare system register reads", "1 February 2022 to 30 June 2024"),
    (F_DASH, "Ridership Data (%s)" % K.PEOPLE["ridership"], "ridership dashboard", "January 2025 to December 2026"),
    (F_COMPLEX, "Ridership Data (%s)" % K.PEOPLE["ridership"], "back office complex table", "as of 11 January 2027"),
    (F_CENSUS, "Station Planning (%s)" % K.PEOPLE["counts"], "count programme", "October 2026"),
    (F_COUNTS, "Station Planning (%s)" % K.PEOPLE["counts"], "count programme", "autumn 2026"),
    (F_CAP, "Station Planning (%s)" % K.PEOPLE["planning_manager"], "platform capacity register", "as of 11 January 2027"),
    (F_TT, "Service Planning (%s)" % K.PEOPLE["service_planner"], "scheduling system", "autumn 2026 and spring 2027 picks"),
    (F_FLEET, "Car Equipment", "fleet register", "as of 11 January 2027"),
    (F_CHANGES, "Service Planning (%s)" % K.PEOPLE["service_planner"], "service change register", "2024 to 2027"),
    (F_INV, "Station Engineering", "engineering inventory", "as of 11 January 2027"),
    (F_ASSET, "Elevator and Escalator Maintenance", "maintenance system asset register", "as of 11 January 2027"),
    (F_CTRL, "Elevator and Escalator Maintenance", "escalator controllers", "1 September to 31 December 2026"),
    (F_WO, "Elevator and Escalator Maintenance", "maintenance system work orders", "1 September to 31 December 2026"),
    (F_MON, "Elevator and Escalator Maintenance", "elevator remote monitoring", "1 September to 31 December 2026"),
]
DICT_FIELDS = [
    (F_FARE, "hour_utc", "Hour starting, UTC."),
    (F_FARE, "complex_id", "Back office complex id (complex_reference.csv)."),
    (F_FARE, "entries", "Fare validations at the complex's arrays in the hour, every payment type."),
    (F_FARE, "paid_contactless", "Entries paid by contactless bank card or phone."),
    (F_FARE, "paid_card", "Entries paid by Vellmouth Card."),
    (F_FARE, "free_transfer", "Bus-to-subway transfers."),
    (F_FARE, "staff", "Staff passes."),
    (F_LEGACY, "svc_date", "Service day, local."),
    (F_LEGACY, "period_end", "Local time the register period ended; the period ending 04:00 is the four hours after midnight."),
    (F_LEGACY, "stn", "Legacy station code."),
    (F_LEGACY, "stn_name", "Station name."),
    (F_LEGACY, "entries", "Entries registered in the period."),
    (F_COMPLEX, "legacy_codes", "Legacy station codes that make up the complex."),
    (F_CAP, "rated_capacity", "People, at one square metre each over the usable area."),
    (F_CAP, "effective_from", "First day the rating applies."),
    (F_CAP, "effective_to", "Last day the rating applies; blank while current."),
    (F_CAP, "basis", "How the rating was made."),
    (F_TT, "plan", "AUT26 autumn 2026 pick (also the winter 2026-27 pick); SPR27 spring 2027 pick."),
    (F_TT, "arrive", "Scheduled arrival, local time."),
    (F_TT, "cars", "Cars per train."),
    (F_FLEET, "max_load_per_car", "People per car, seated and standing, at the design maximum."),
    (F_FLEET, "planning_load_per_car", "People per car used for service planning loads."),
    (F_INV, "unit", "Element id: complex id, then S stair, E escalator, L elevator."),
    (F_INV, "from_area", "Lower or outer end."),
    (F_INV, "to_area", "Other end."),
    (F_INV, "width", "Stair width; escalator step width."),
    (F_INV, "width_unit", "in or mm."),
    (F_INV, "direction", "Escalator running direction."),
    (F_INV, "speed_mps", "Escalator design speed."),
    (F_INV, "surveyed", "Condition survey."),
    (F_ASSET, "step_width_mm", "Step width."),
    (F_ASSET, "speed_mps", "Speed the controller is set to."),
    (F_CTRL, "event_utc", "UTC."),
    (F_CTRL, "event", "RUN UP, RUN DOWN, STOP, FAULT."),
    (F_WO, "type", "CM corrective; PM planned maintenance, the car on inspection operation for the work."),
    (F_WO, "out_of_service_at", "UTC."),
    (F_WO, "in_service_at", "UTC."),
    (F_MON, "event_utc", "UTC."),
    (F_MON, "mode", "The car's controller mode as the monitoring receives it, from that time: NORMAL, FAULT, INDEPENDENT, SHUTDOWN."),
]


def write_dictionary(out):
    path = os.path.join(out, F_DICT)
    wb = Workbook()
    ws = wb.active
    ws.title = "Files"
    _sheet(ws, ["File", "Owner", "Source", "Period"], DICT_FILES, [52, 40, 40, 34])
    ws2 = wb.create_sheet("Fields")
    _sheet(ws2, ["File", "Field", "Description"], DICT_FIELDS, [52, 22, 90])
    wb.save(path)
    PP.normalize_ooxml(path, K.PEOPLE["ridership"], K.PEOPLE["ridership"], "2024-07-15T13:10:00Z",
                       "2027-01-11T14:22:00Z", excel=True, company=K.ORG)


# ---------------------------------------------------------------------------------------------------- the rest
def write_fare(out, fare):
    path = os.path.join(out, F_FARE)
    ts = pa.array([r[0] for r in fare], type=pa.timestamp("s", tz="UTC"))
    cols = [ts] + [pa.array([r[i] for r in fare], type=pa.int32()) for i in range(1, 7)]
    t = pa.Table.from_arrays(cols, names=R.FARE_COLS)
    pq.write_table(t, path, compression="snappy", row_group_size=200000)
    return len(fare)


def build_pack(out):
    os.makedirs(out, exist_ok=True)
    p = lambda f: os.path.join(out, f)
    L = R.Levels()
    n = {}
    fare = R.fare_hourly(L)
    n[F_FARE] = write_fare(out, fare)
    n[F_LEGACY] = K.write_csv(p(F_LEGACY), R.LEGACY_COLS, R.legacy_rows(L))
    n[F_DASH] = write_dashboard(out, fare)
    n[F_COMPLEX] = K.write_csv(p(F_COMPLEX), R.COMPLEX_COLS, R.complex_rows())
    n[F_CENSUS] = write_census(out, L)
    n[F_COUNTS] = write_counts(out, L)
    n[F_CAP] = K.write_csv(p(F_CAP), R.CAP_COLS, R.capacity_rows())
    n[F_TT] = K.write_csv(p(F_TT), R.TT_COLS, R.timetable_rows())
    n[F_FLEET] = K.write_csv(p(F_FLEET), R.FLEET_COLS, R.fleet_rows())
    n[F_CHANGES] = K.write_csv(p(F_CHANGES), R.CHANGE_COLS, R.change_rows())
    n[F_INV] = K.write_csv(p(F_INV), R.INV_COLS, R.inventory_rows())
    n[F_ASSET] = K.write_csv(p(F_ASSET), R.ASSET_COLS, R.asset_rows())
    ctrl = [(_ts(_round_s(t)), u, e) for (t, u, e) in R.controller_rows()]
    n[F_CTRL] = K.write_csv(p(F_CTRL), R.CTRL_COLS, ctrl)
    wos, mon = R.elevator_records()
    n[F_WO] = K.write_csv(p(F_WO), R.WO_COLS, [w[:6] + (_ts(_round_s(w[6])), _ts(_round_s(w[7]))) for w in wos])
    n[F_MON] = K.write_csv(p(F_MON), R.MON_COLS, [(_ts(_round_s(t)), u, m) for (t, u, m) in mon])
    write_dictionary(out)
    DOC.write_all(out)
    return n
