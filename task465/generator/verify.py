"""Independent verifier for task465: reads only the shipped folder and the goldens, on its own code path (pandas, numpy,
zoneinfo), and recomputes each division and band's households and survey shares, the pilot's tables and odds ratios,
the fused eligible households and the caseworkers, and the two asks' household counts, against generator/figures.json,
the golden workbook and the golden table. Imports nothing from the generator.

    python3 task465/generator/verify.py
"""
import json
import math
import os
import sys
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
CASELOAD = 400


def f(target, stem):
    hit = [x for x in os.listdir(target) if x.startswith(stem)]
    assert len(hit) == 1, (stem, hit)
    return os.path.join(target, hit[0])


def fuse(p, q, psi):
    """Share with both traits in the 2x2 table with margins p and q and odds ratio psi."""
    a = psi - 1.0
    b = -(1.0 + a * (p + q))
    c = psi * p * q
    return (-b - math.sqrt(b * b - 4 * a * c)) / (2 * a)


def shares(target):
    reg = pd.read_csv(f(target, "callowmere_service_divisions"), dtype=str)
    div = dict(zip(reg["COUNTY_FIPS"], reg["SERVICE_DIVISION"]))
    h = pd.read_csv(f(target, "heas_2025_household_puf"), dtype={"COUNTY": str})
    h["div"] = h["COUNTY"].map(div)
    h["band"] = np.where(h["HHINC"] < 25000, 0, np.where(h["HHINC"] < 50000, 1, 2))
    h["cost"] = np.where(h["ELEC_BASIS"] == "M", 12, 1) * h["ELEC_AMT"] + h["GAS_AMT"] + h["OTHER_AMT"]
    h["hi"] = (h["cost"] / h["HHINC"] > 0.06).astype(float)
    h = h[h["div"].notna() & (h["band"] < 2)]
    p = (h["hi"] * h["WT_FINAL"]).groupby([h["div"], h["band"]]).sum() / h.groupby(["div", "band"])["WT_FINAL"].sum()
    s = pd.read_csv(f(target, "shs_2025_household_file"))
    s["div"] = s["CNTYFIPS"].map(lambda v: div.get("%03d" % v))
    s["band"] = s["INCCAT"].map(lambda c: 0 if c <= 3 else 1 if c <= 5 else 2)
    s = s[s["div"].notna() & (s["band"] < 2) & s["MEDEQUIP"].isin([1, 2])]
    s["eq"] = (s["MEDEQUIP"] == 1).astype(float)
    q = (s["eq"] * s["HH_WT"]).groupby([s["div"], s["band"]]).sum() / s.groupby(["div", "band"])["HH_WT"].sum()
    e = pd.read_excel(f(target, "household_estimates_county_income_2025"), sheet_name="Households by income", header=3,
                      dtype={"FIPS": str})
    e = e[e["FIPS"].isin(div)]
    e["div"] = e["FIPS"].map(div)
    n = pd.concat([e.groupby("div")["Under $25,000"].sum().rename(0), e.groupby("div")["$25,000 to $49,999"].sum()
                   .rename(1)], axis=1).stack()
    n.index.names = ["div", "band"]
    return p, q, n, div


def pilot(target):
    s = pd.read_csv(f(target, "hema_screening_pilot_2025"))
    s = s[s["STATUS"] == "Complete"].sort_values(["RECEIVED", "SCREEN_ID"]).groupby("ADDRESS_KEY").tail(1)
    s = s[s["MED_EQUIPMENT"].isin(["Yes", "No"])]
    s["band"] = np.where(s["HH_INCOME"] < 25000, 0, 1)
    s["x"] = (s["ELEC_12M"] + s["OTHER_FUELS_12M"]) / s["HH_INCOME"] > 0.06
    s["y"] = s["MED_EQUIPMENT"] == "Yes"
    psi = {}
    for b, t in s.groupby("band"):
        a = ((t.x) & (t.y)).sum()
        bb = ((t.x) & (~t.y)).sum()
        c = ((~t.x) & (t.y)).sum()
        d = ((~t.x) & (~t.y)).sum()
        psi[b] = a * d / (bb * c)
    return psi, s


def households_on_rate(target):
    e = pd.read_csv(f(target, "medical_baseline_enrolment_2026"), dtype=str, keep_default_na=False)
    root = {a: a for a in e["ACCOUNT_NO"]}
    for a, prior in zip(e["ACCOUNT_NO"], e["PRIOR_ACCOUNT_NO"]):
        if prior:
            root.setdefault(prior, prior)

    def find(a):
        while root[a] != a:
            root[a] = root[root[a]]
            a = root[a]
        return a
    for a, prior in zip(e["ACCOUNT_NO"], e["PRIOR_ACCOUNT_NO"]):
        if prior:
            ra, rb = find(a), find(prior)
            if ra != rb:
                root[max(ra, rb)] = min(ra, rb)
    return e, find


def main(target, golden, fig_path=None):
    F = json.load(open(fig_path or os.path.join(HERE, "figures.json")))
    bad = []

    def check(label, cond, detail=""):
        print("%s  verify: %s%s" % ("PASS" if cond else "FAIL", label, ("  [" + str(detail) + "]") if detail != "" else ""))
        if not cond:
            bad.append(label)

    p, q, n, div = shares(target)
    psi, scr = pilot(target)
    check("the pilot's odds ratios inside each band", all(round(psi[b], 2) == F["psi"][str(b)] for b in (0, 1)),
          {b: round(v, 3) for b, v in psi.items()})
    cells = {(d, b): n[(d, b)] * fuse(p[(d, b)], q[(d, b)], psi[b]) for d, b in n.index}
    divs = sorted({d for d, _ in cells})
    tot = {d: cells[(d, 0)] + cells[(d, 1)] for d in divs}
    staff = {d: int(math.ceil(tot[d] / CASELOAD)) for d in divs}
    check("each division's caseworkers and the total", staff == F["call"]["by_division"]
          and sum(staff.values()) == F["call"]["total"], staff)
    r = lambda v: int(math.floor(v / 100 + 0.5) * 100)
    check("each division's eligible households to the nearest hundred", {d: r(tot[d]) for d in divs} == F["eligible"])
    check("each cell's eligible households to the nearest hundred",
          all(r(cells[(d, b)]) == F["cells"]["%s|%d" % (d, b)] for d, b in cells))
    check("each cell's households, the demographer's county rows summed",
          all(int(n[(d, b)]) == F["households"]["%s|%d" % (d, b)] for d, b in cells))
    indep = {d: int(math.ceil((n[(d, 0)] * p[(d, 0)] * q[(d, 0)] + n[(d, 1)] * p[(d, 1)] * q[(d, 1)]) / CASELOAD))
             for d in divs}
    check("the two surveys' shares multiplied staff %d" % sum(indep.values()),
          sum(indep.values()) == F["rivals"]["independence"]["total"])

    e, find = households_on_rate(target)
    e["start"] = pd.to_datetime(e["START_DATE"])
    e["end"] = pd.to_datetime(e["END_DATE"].replace("", None))
    cut = pd.Timestamp("2026-10-31")
    act = e[(e["start"] <= cut) & (e["end"].isna() | (e["end"] >= cut))].copy()
    reg = pd.read_csv(f(target, "callowmere_service_divisions"), dtype=str)
    cdiv = dict(zip(reg["COUNTY"], reg["SERVICE_DIVISION"]))
    act["hh"] = act["ACCOUNT_NO"].map(find)
    act["div"] = act["SERVICE_COUNTY"].map(cdiv)
    a = act.sort_values("start").groupby("hh").tail(1).groupby("div").size().to_dict()
    check("Ask A: households on the medical baseline rate on 31 October 2026", a == F["baseline"], a)
    o = pd.read_csv(f(target, "battery_placement_work_orders_2026"), dtype=str, keep_default_na=False)
    o = o[o["STATUS"] == "Completed"].copy()
    local = pd.to_datetime(o["COMPLETED_UTC"], utc=True).dt.tz_convert(ZoneInfo("America/Chicago"))
    o["day"] = local.dt.date.astype(str)
    o = o[(o["day"] >= "2026-01-01") & (o["day"] <= "2026-10-31")]
    o["hh"] = o["ACCOUNT_NO"].map(lambda x: find(x) if x in set(e["ACCOUNT_NO"]) else x)
    o["div"] = o["SERVICE_COUNTY"].map(cdiv)
    bcount = o.groupby("hh")["div"].last().value_counts().to_dict()
    check("Ask B: households given a backup battery from January to October 2026", bcount == F["battery"], bcount)

    import openpyxl
    ws = openpyxl.load_workbook(os.path.join(golden, "hema_caseworkers_2027.xlsx"))["Caseworkers 2027"]
    rows = {ws.cell(row=i, column=1).value: [ws.cell(row=i, column=j).value for j in range(2, 6)] for i in range(5, 15)}
    check("the golden workbook", all(rows[d] == [r(tot[d]), staff[d], a[d], bcount[d]] for d in divs))
    t = pd.read_csv(os.path.join(golden, "hema_eligible_households_2025.csv"))
    band = {"Under $25,000": 0, "$25,000 to $49,999": 1}
    check("the golden table", len(t) == 18 and all(
        row.ELIGIBLE_HOUSEHOLDS_2025 == r(cells[(row.SERVICE_DIVISION, band[row.INCOME_BAND])])
        for row in t.itertuples()))
    if bad:
        raise SystemExit("verify failed: %s" % bad)


if __name__ == "__main__":
    main(os.path.join(TASK, "target"), os.path.join(TASK, "golden"))
