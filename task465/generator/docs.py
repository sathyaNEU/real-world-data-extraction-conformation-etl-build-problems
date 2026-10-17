"""The pack's documents: the energy survey's technical notes, the health survey's codebook, the screening pilot's
operations report, the 2027 HEMA plan, the folder's data dictionary and its file register."""
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import letter                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import inch                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import pack as P    # noqa: E402
import world as W   # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("LiberationSerif", "Regular"), ("LiberationSerif-Bold", "Bold"), ("LiberationSans", "Regular"),
              ("LiberationSans-Bold", "Bold")):
    try:
        pdfmetrics.getFont(nm)
    except KeyError:
        pdfmetrics.registerFont(TTFont(nm, FONT_DIR + nm.split("-")[0] + "-" + f + ".ttf"))
INK = colors.HexColor("#1b1d1f")
GREY = colors.HexColor("#5b6067")
TEAL = colors.HexColor("#1d5c63")
NAVY = colors.HexColor("#1f3a5f")
PLUM = colors.HexColor("#5a2d52")
BODY = ParagraphStyle("body", fontName="LiberationSerif", fontSize=10.5, leading=14.2, textColor=INK, spaceAfter=6)
NUM = ParagraphStyle("num", parent=BODY, leftIndent=28, firstLineIndent=-28)
CELL = ParagraphStyle("cell", fontName="LiberationSans", fontSize=8.8, leading=11, textColor=INK)
CELLB = ParagraphStyle("cellb", parent=CELL, fontName="LiberationSans-Bold")
SMALL = ParagraphStyle("small", fontName="LiberationSans", fontSize=8.5, leading=11, textColor=GREY)


def styles(accent):
    h1 = ParagraphStyle("h1", fontName="LiberationSans-Bold", fontSize=14, leading=18, textColor=accent, spaceAfter=2)
    h1s = ParagraphStyle("h1s", fontName="LiberationSans", fontSize=9.3, leading=12.5, textColor=GREY, spaceAfter=10)
    h2 = ParagraphStyle("h2", fontName="LiberationSans-Bold", fontSize=10.5, leading=14, textColor=INK, spaceBefore=7,
                        spaceAfter=3)
    return h1, h1s, h2


def _doc(path, title, author, footer):
    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFont("LiberationSans", 7.5)
        canvas.setFillColor(GREY)
        canvas.drawString(0.9 * inch, 0.6 * inch, footer)
        canvas.drawRightString(letter[0] - 0.9 * inch, 0.6 * inch, "Page %d" % doc.page)
        canvas.restoreState()
    d = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.9 * inch, rightMargin=0.9 * inch, topMargin=0.8 * inch,
                          bottomMargin=0.9 * inch, title=title, author=author, subject="", creator="", producer="")
    return d, frame


def _table(rows, widths, head="#e6eef0"):
    data = [[Paragraph(str(c), CELLB if i == 0 else CELL) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa0a6")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(head)),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 4),
                           ("RIGHTPADDING", (0, 0), (-1, -1), 4)]))
    return t


def fmt(v):
    return "{:,}".format(int(v))


# ---------------------------------------------------------------- the energy survey's technical notes
def write_heas_notes(path, n):
    H1, H1S, H2 = styles(TEAL)
    doc, frame = _doc(path, "Household Energy Affordability Survey 2025: Technical Notes", W.ENERGY_OFFICE,
                      "%s  |  Household Energy Affordability Survey 2025  |  Technical notes for the public-use "
                      "household file  |  November 2025" % W.ENERGY_OFFICE)
    s = [Paragraph("Household Energy Affordability Survey 2025", H1),
         Paragraph("Technical notes for the public-use household file. %s, Research and Planning Section. Released "
                   "November 20, 2025." % W.ENERGY_OFFICE, H1S),
         Paragraph("1. About the survey", H2),
         Paragraph("The Household Energy Affordability Survey measures what households in the state pay for home "
                   "energy and how those costs compare with household income. The 2025 survey was in the field from "
                   "February 10 to June 27, 2025, and asked about electricity, natural gas and other heating fuels, "
                   "the home and its heating, and household income.", BODY),
         Paragraph("2. Sample", H2),
         Paragraph("The sample was drawn from the state's residential address frame, stratified by county, with at "
                   "least 70 completed interviews in each county and the rest allocated in proportion to the county's "
                   "households. Households answered on the web, by telephone or on a paper form (INTMODE W, T and P). "
                   "%s households completed the survey; the public-use file holds one record for each." % fmt(n), BODY),
         Paragraph("3. Weights", H2),
         Paragraph("WT_BASE is the inverse of the household's probability of selection, adjusted for nonresponse "
                   "within county and response mode. WT_FINAL calibrates WT_BASE to the State Demographer's estimates "
                   "of households by county and household income band for July 1, 2025 (under $25,000; $25,000 to "
                   "$49,999; $50,000 to $74,999; $75,000 and over), so that the weighted households of every county "
                   "and band equal those estimates. The office's published estimates use WT_FINAL.", BODY),
         Paragraph("4. Energy costs", H2),
         Paragraph("ELEC_AMT is what the household paid for electricity over the past 12 months. Respondents who "
                   "could not give a 12-month amount gave a typical monthly bill instead; for them ELEC_BASIS is M and "
                   "the annual electricity cost is 12 times ELEC_AMT. For everyone else ELEC_BASIS is A. GAS_AMT is "
                   "the 12-month cost of utility natural gas and OTHER_AMT the 12-month cost of bottled gas, fuel oil, "
                   "kerosene, wood and any other heating fuel. A household's annual home energy cost is its annual "
                   "electricity cost plus GAS_AMT plus OTHER_AMT. Amounts are whole dollars and include anything an "
                   "assistance program paid on the household's behalf.", BODY),
         Paragraph("The office reports a household's energy burden as its annual home energy cost as a share of its "
                   "household income.", BODY),
         Paragraph("5. Income, home and geography", H2),
         Paragraph("HHINC is total household money income over the 12 months before the interview, in dollars, "
                   "rounded to the nearest $100 and top-coded at $248,000; missing income was imputed before release. "
                   "TENURE is 1 for a home owned and 2 for a home rented. HEATFUEL is the main heating fuel: 1 utility "
                   "gas, 2 electricity, 3 bottled, tank or LP gas, 4 fuel oil or kerosene, 5 wood, 6 other or none. "
                   "COUNTY is the three-digit county code; the file carries no geography below the county.", BODY),
         Paragraph("6. Use of the file", H2),
         Paragraph("Estimates for areas smaller than a county, or for counties and income bands with few completed "
                   "interviews, carry wide margins of error. Questions about the file go to the Research and Planning "
                   "Section.", BODY)]
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------- the health survey's codebook
def write_shs_codebook(path):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Variables"
    ws["A1"] = "State Health Survey 2025: household file codebook"
    ws["A1"].font = Font(name="Calibri", size=13, bold=True)
    ws["A2"] = "%s, Center for Health Statistics. Household file released May 12, 2026." % W.HEALTH_DEPT
    ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="595959")
    head = ["Variable", "Description", "Values", "Notes"]
    for j, h in enumerate(head, start=1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(name="Calibri", size=10, bold=True)
        c.fill = PatternFill("solid", fgColor="E4DFEC")
    cats = "; ".join("%d %s" % (k, lab) for k, lab, _ in W.INCCAT)
    rows = [
        ("SEQNO", "Household sequence number", "", ""),
        ("CNTYFIPS", "County of residence", "County code (numeric)", ""),
        ("PHD", "Public health district", "1 to 7", "The department's districts; they do not follow county "
         "groupings used by other agencies."),
        ("INCCAT", "Household income, last 12 months", cats, "Missing income was imputed before release."),
        ("ADULTS", "Adults aged 18 and over in the household", "1 to 4 (4 = 4 or more)", ""),
        ("CHILDREN", "Children under 18 in the household", "0 to 3 (3 = 3 or more)", ""),
        ("AGEGRP", "Respondent's age", "1 18 to 24; 2 25 to 34; 3 35 to 44; 4 45 to 54; 5 55 to 64; 6 65 or older",
         ""),
        ("GENHLTH", "Respondent's general health", "1 Excellent; 2 Very good; 3 Good; 4 Fair; 5 Poor; 7 Don't know; "
         "9 Refused", ""),
        ("MEDEQUIP", "Does anyone living in this household use electrically powered medical equipment at home, such "
         "as an oxygen concentrator, a home dialysis machine, a ventilator, a powered wheelchair or a feeding pump?",
         "1 Yes; 2 No; 7 Don't know; 9 Refused", ""),
        ("HH_WT", "Household weight", "", "See Notes."),
    ]
    for i, r in enumerate(rows, start=5):
        for j, v in enumerate(r, start=1):
            c = ws.cell(row=i, column=j, value=v or None)
            c.font = Font(name="Calibri", size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCD", (12, 58, 52, 40)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    ws2 = wb.create_sheet("Notes")
    notes = [
        "The household file holds one record for each household that completed the 2025 survey.",
        "HH_WT is the household weight: the inverse of the household's probability of selection, adjusted for "
        "nonresponse and calibrated to the State Demographer's estimates of households by county and household income "
        "band for July 1, 2025 (INCCAT 1 to 3, 4 to 5, 6, and 7 to 8). Use HH_WT for estimates of households.",
        "Estimates of percentages exclude don't know (7) and refused (9) responses.",
        "County codes are those of the state's county list; the file carries no geography below the county.",
    ]
    ws2["A1"] = "Notes"
    ws2["A1"].font = Font(name="Calibri", size=12, bold=True)
    for i, t in enumerate(notes, start=3):
        c = ws2.cell(row=i, column=1, value=t)
        c.font = Font(name="Calibri", size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws2.column_dimensions["A"].width = 110
    wb.save(path)


# ---------------------------------------------------------------- the screening pilot's operations report
def write_pilot_report(path, screens):
    H1, H1S, H2 = styles(NAVY)
    doc, frame = _doc(path, "HEMA Screening Pilot 2025: Operations Report", W.ORG,
                      "%s  |  Customer Programmes  |  HEMA screening pilot 2025: operations report  |  September 2025"
                      % W.ORG)
    import sim as S
    tab = [["Division", "Mode", "Screens returned", "Ended at the income question", "Complete", "Partial"]]
    for d, m, v in S.pilot_counts(screens):
        tab.append([d, m] + [fmt(n) for n in v])
    s = [Paragraph("HEMA screening pilot 2025: operations report", H1),
         Paragraph("%s, Customer Programmes. Prepared by %s, %s. September 26, 2025." % (
             W.ORG, W.PEOPLE["manager"], W.TITLES["manager"]), H1S),
         Paragraph("1. Purpose", H2),
         Paragraph("Before the 2027 programme plan, Customer Programmes tested the Home Energy Medical Assistance "
                   "screener, the mail and telephone workflow and the hand-off to caseworkers in the North Metro and "
                   "South Metro divisions.", BODY),
         Paragraph("2. Addresses and contact", H2),
         Paragraph("The billing system drew 10,800 active residential service addresses at random in each of the two "
                   "divisions, 21,600 in all. Each address was mailed the screener on March 3, 2025, with a reminder "
                   "card on March 24. From May 5 to August 29, 2025 the call centre telephoned addresses that had not "
                   "returned a screen, and some that had.", BODY),
         Paragraph("3. The screener", H2),
         Paragraph("Q1 asks the household's income over the last 12 months; a household reporting $50,000 or more is "
                   "thanked and the screen ends there. Q2 asks what the household paid for electricity over the last 12 "
                   "months, Q3 what it paid for natural gas, propane, fuel oil, wood or other heating fuels over the "
                   "same months, and Q4 whether anyone living in the home uses electrically powered medical equipment "
                   "at home (yes, no, don't know).", BODY),
         Paragraph("4. Returns", H2),
         _table(tab, [1.15 * inch, 0.7 * inch, 1.05 * inch, 1.45 * inch, 0.9 * inch, 0.8 * inch]),
         Spacer(1, 8),
         Paragraph("5. The screen records", H2),
         Paragraph("Each returned screen is one record in the pilot file. An address that returned a screen by mail "
                   "and answered again when the call centre telephoned has two records under the same address key; the "
                   "programme takes the latest complete screen as the address's answer. A partial screen stopped before "
                   "Q4.", BODY),
         Paragraph("6. What the pilot taught us", H2),
         Paragraph("A telephone screen took 6.4 minutes on average and a returned paper screen 3 minutes to key. Mail "
                   "returns peaked in the second week after each mailing. Several callers read Q3 as covering only "
                   "natural gas; the 2027 screener names propane and fuel oil first. The call centre needs two more "
                   "agents during a mailing month at the scale of a territory-wide screen.", BODY)]
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------- the 2027 HEMA plan
def write_plan(path):
    H1, H1S, H2 = styles(PLUM)
    doc, frame = _doc(path, "Home Energy Medical Assistance: 2027 Programme Plan (draft)", W.ORG,
                      "%s  |  Home Energy Medical Assistance  |  2027 programme plan, draft for filing  |  November 2026"
                      % W.ORG)
    s = [Paragraph("Home Energy Medical Assistance: 2027 programme plan", H1),
         Paragraph("Draft for filing with the %s by December 1, 2026. Prepared by %s, %s; reviewed by %s, %s. "
                   "November 12, 2026." % (W.REGULATOR, W.PEOPLE["manager"], W.TITLES["manager"], W.PEOPLE["vp"],
                                           W.TITLES["vp"]), H1S),
         Paragraph("1. Purpose", H2),
         Paragraph("Home Energy Medical Assistance (HEMA) helps households with modest incomes in which someone relies "
                   "on electrically powered medical equipment at home. Oxygen concentrators, home dialysis machines, "
                   "ventilators and similar equipment run for hours every day and add to a household's electricity "
                   "bill, and an interruption in service can endanger the person who relies on them.", BODY),
         Paragraph("2. Eligibility", H2),
         Paragraph("A household is eligible for HEMA when (a) its household income is under $50,000, (b) its annual "
                   "home energy costs (electricity, natural gas and other home heating fuels) exceed 6 per cent of its "
                   "household income, and (c) someone who lives in the home relies on electrically powered medical "
                   "equipment at home.", BODY),
         Paragraph("3. Services", H2),
         Paragraph("HEMA caseworkers enrol eligible households in the medical baseline rate, arrange payment plans "
                   "and budget billing, refer households to energy assistance and weatherization programs, and arrange "
                   "a backup battery for customers who rely on life-support equipment (placements began in December "
                   "2025).", BODY),
         Paragraph("4. Staffing standard", H2),
         Paragraph("HEMA caseworkers are staffed by service division: one caseworker for every 400 eligible households "
                   "in a service division, or part of that number. The service divisions are the nine divisions of the "
                   "company's division register.", BODY),
         Paragraph("5. Data annex: eligible households", H2),
         Paragraph("The staffing rests on the number of eligible households in each service division and eligible "
                   "income band (under $25,000; $25,000 to $49,999) in 2025, the latest survey year, each estimated "
                   "for the division and band as a whole and filed with the plan as one table by division and band. The Commission checks such tables against the state's "
                   "official figures, so in each division and band the table reproduces the State Demographer's "
                   "household estimate for July 1, 2025, the %s's %s estimate of the share of households whose annual "
                   "home energy costs exceed 6 per cent of income, and the %s's %s estimate of the share of households "
                   "in which someone relies on electrically powered medical equipment."
                   % (W.ENERGY_OFFICE, W.SURVEY_A, W.HEALTH_DEPT, W.SURVEY_B), BODY),
         Paragraph("6. Budget", H2),
         Paragraph("The 2027 budget carries each caseworker at $78,400 in salary and benefits, backup batteries at "
                   "$1,150 a unit installed, and outreach at $236,000. The caseworker line follows the staffing in "
                   "section 4.", BODY),
         Paragraph("7. Filing and timetable", H2),
         Paragraph("The plan is filed with the %s by December 1, 2026. The company expects the Commission's decision "
                   "in March 2027 and staffs the 2027 caseworker positions from April 1, 2027." % W.REGULATOR, BODY)]
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------- the folder's dictionary and register
DICT = [
    ("heas", "RESPID", "Household record number"),
    ("heas", "COUNTY", "County code, three digits"),
    ("heas", "INTMODE", "Response mode: W web, T telephone, P paper"),
    ("heas", "TENURE", "1 owned, 2 rented"),
    ("heas", "HEATFUEL", "Main heating fuel (see the technical notes)"),
    ("heas", "HHINC", "Household income over the 12 months before the interview, dollars"),
    ("heas", "ELEC_AMT", "Electricity cost, dollars (see ELEC_BASIS)"),
    ("heas", "ELEC_BASIS", "A: ELEC_AMT covers the past 12 months; M: ELEC_AMT is a typical monthly bill"),
    ("heas", "GAS_AMT", "Utility natural gas, past 12 months, dollars"),
    ("heas", "OTHER_AMT", "Bottled gas, fuel oil, kerosene, wood and other fuels, past 12 months, dollars"),
    ("heas", "WT_BASE", "Base weight"),
    ("heas", "WT_FINAL", "Final weight"),
    ("shs", "(all fields)", "See shs_2025_household_codebook.xlsx"),
    ("pilot", "SCREEN_ID", "Screen record number"),
    ("pilot", "ADDRESS_KEY", "Service address key"),
    ("pilot", "DIVISION", "Service division of the address"),
    ("pilot", "MODE", "Mail or Phone"),
    ("pilot", "RECEIVED", "Date the screen was received or taken"),
    ("pilot", "STATUS", "Complete; Ended - income (income of $50,000 or more); Partial (stopped before the equipment "
                        "question)"),
    ("pilot", "HH_INCOME", "Household income over the last 12 months, dollars"),
    ("pilot", "ELEC_12M", "Electricity paid over the last 12 months, dollars"),
    ("pilot", "OTHER_FUELS_12M", "Natural gas, propane, fuel oil, wood and other heating fuels paid over the last 12 "
                                 "months, dollars"),
    ("pilot", "MED_EQUIPMENT", "Anyone living in the home uses electrically powered medical equipment at home: Yes, "
                               "No, Don't know"),
    ("estimates", "(sheet Households by income)", "Households by county and household income band, July 1, 2025, with "
                                                 "planning region and state rows"),
    ("divisions", "COUNTY_FIPS", "County code, three digits"),
    ("divisions", "COUNTY", "County"),
    ("divisions", "SERVICE_DIVISION", "Callowmere service division serving the county"),
    ("baseline", "ENROL_ID", "Medical baseline enrolment spell"),
    ("baseline", "ACCOUNT_NO", "Callowmere account billed at the medical baseline rate"),
    ("baseline", "PRIOR_ACCOUNT_NO", "The customer's previous account, where the medical baseline was transferred from "
                                     "it on a move"),
    ("baseline", "PREMISE_ID", "Service premise of the account"),
    ("baseline", "SERVICE_COUNTY", "County of the premise"),
    ("baseline", "DIVISION", "Service division recorded on the account when it was set up"),
    ("baseline", "START_DATE", "First day billed at the medical baseline rate"),
    ("baseline", "END_DATE", "Last day billed at the medical baseline rate; blank while the spell is open"),
    ("baseline", "END_REASON", "Why the spell ended"),
    ("baseline", "QUALIFYING_EQUIPMENT", "Equipment named on the medical certification"),
    ("baseline", "LIFE_SUPPORT", "Y where the certification names life-support equipment"),
    ("batteries", "WORK_ORDER", "Field work order"),
    ("batteries", "ACCOUNT_NO", "Account served"),
    ("batteries", "PREMISE_ID", "Premise visited"),
    ("batteries", "SERVICE_COUNTY", "County of the premise"),
    ("batteries", "SCHEDULED", "Date the visit was booked for (Central Time, the company's operating day)"),
    ("batteries", "STATUS", "Completed, Cancelled or No access"),
    ("batteries", "COMPLETED_UTC", "Completion time from the crew's handheld, UTC"),
    ("batteries", "UNIT_SERIAL", "Battery unit installed"),
    ("batteries", "REMOVED_SERIAL", "Battery unit taken out, where one was"),
    ("batteries", "CREW", "Field crew"),
    ("hvi", "COUNTY_FIPS", "County code, three digits"),
    ("hvi", "COUNTY", "County"),
    ("hvi", "HVI_SCORE", "Heat vulnerability index (state mean 50, standard deviation 15)"),
    ("hvi", "PCT_AGE_65_PLUS", "Residents aged 65 and over, per cent"),
    ("hvi", "PCT_65_PLUS_LIVING_ALONE", "Residents aged 65 and over living alone, per cent"),
    ("hvi", "PCT_NO_CENTRAL_AC", "Homes without central air conditioning, per cent"),
    ("hvi", "PCT_IMPERVIOUS", "Impervious surface, per cent of land"),
    ("hvi", "HVI_RANK", "Rank of the county's score, 1 most vulnerable"),
    ("liheap", "COUNTY_FIPS", "County code, three digits"),
    ("liheap", "COUNTY", "County"),
    ("liheap", "MONTH", "Month the application was received"),
    ("liheap", "APPLICATIONS_RECEIVED", "Energy assistance applications received"),
    ("liheap", "APPLICATIONS_APPROVED", "Applications approved"),
    ("liheap", "CRISIS_APPLICATIONS", "Applications for crisis assistance (shut-off notice or empty tank)"),
    ("liheap", "AVERAGE_BENEFIT_USD", "Average approved benefit, dollars"),
]


def write_dictionary(path):
    rows = [[W.FILES[k], f, d] for k, f, d in DICT]
    return P.write_csv(path, ["SOURCE_FILE", "COLUMN", "MEANING"], rows)


REGISTER = [
    ("heas", W.ENERGY_OFFICE, "Household Energy Affordability Survey 2025, public-use household file", "2025-11-20"),
    ("heas_methods", W.ENERGY_OFFICE, "Technical notes for the public-use household file", "2025-11-20"),
    ("shs", W.HEALTH_DEPT, "State Health Survey 2025, household file", "2026-05-12"),
    ("shs_codebook", W.HEALTH_DEPT, "Codebook for the household file", "2026-05-12"),
    ("pilot", "Callowmere Customer Programmes", "HEMA screening pilot 2025, screen records", "2025-09-26"),
    ("pilot_report", "Callowmere Customer Programmes", "HEMA screening pilot 2025, operations report", "2025-09-26"),
    ("estimates", W.DEMOGRAPHER, "County households by household income band, July 1, 2025", "2026-03-18"),
    ("divisions", "Callowmere Network Planning", "Division register: the service division of each county served",
     "2026-11-02"),
    ("plan", "Callowmere Customer Programmes", "HEMA 2027 programme plan, draft for filing", "2026-11-12"),
    ("baseline", "Callowmere Customer Information System", "Medical baseline enrolment spells active since January 1, "
     "2020", "2026-11-03"),
    ("batteries", "Callowmere Field Operations", "Backup battery work orders", "2026-11-03"),
    ("hvi", W.HEALTH_DEPT, "Heat vulnerability index by county, 2025", "2025-07-15"),
    ("liheap", W.ENERGY_OFFICE, "Energy assistance applications by county and month, 2025", "2026-01-21"),
    ("dict", "Callowmere Customer Programmes", "Data dictionary for this folder", "2026-11-16"),
]


def write_register(path, counts):
    rows = []
    for k, src, desc, asof in REGISTER:
        rows.append([W.FILES[k], src, desc, asof, counts.get(k, "")])
    return P.write_csv(path, ["FILE_NAME", "PROVIDED_BY", "CONTENTS", "CURRENT_AS_OF", "RECORDS"], rows)
