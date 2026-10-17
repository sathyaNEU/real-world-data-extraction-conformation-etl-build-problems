"""Write the pack's data files: the two surveys' household files, the screening pilot's records, the demographer's
household estimates, the division register, the medical baseline enrolment extract, the battery work orders and the two
neighbouring state files."""
import csv
import datetime as dt

import numpy as np

import sim as S
import world as W


def write_csv(path, header, rows, delimiter=",", lineterminator="\r\n"):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh, delimiter=delimiter, lineterminator=lineterminator)
        wr.writerow(header)
        for r in rows:
            wr.writerow(r)
    return len(rows)


def money(v):
    return "%.2f" % v


def weight(v):
    return "%.4f" % v


def numbers(key, n, lo, hi):
    """Record numbers as a system issues them: increasing, with the gaps left by records not in the extract."""
    rng = np.random.default_rng([S.SEED, 30, key])
    return np.sort(rng.choice(np.arange(lo, hi), size=n, replace=False))


# ---------------------------------------------------------------- the energy office's public-use file
HEAS_COLS = ["RESPID", "COUNTY", "INTMODE", "TENURE", "HEATFUEL", "HHINC", "ELEC_AMT", "ELEC_BASIS", "GAS_AMT",
             "OTHER_AMT", "WT_BASE", "WT_FINAL"]


def heas_rows(x):
    A = x["A"]
    wf = S.final_weights(A["base"], A["county"], A["band"], x["pub"])
    rng = np.random.default_rng([S.SEED, 21])
    order = rng.permutation(len(A["base"]))
    ids = numbers(1, len(order), 10001, 99999)
    rows = []
    for k, i in enumerate(order):
        rows.append(["H25%05d" % ids[k], W.FIPS[A["county"][i]], A["mode"][i], int(A["tenure"][i]),
                     int(A["fuel"][i]), int(A["income"][i]), int(A["elec_amt"][i]), "M" if A["monthly"][i] else "A",
                     int(A["gas"][i]), int(A["other"][i]), weight(A["base"][i]), weight(wf[i])])
    return rows


# ---------------------------------------------------------------- the health department's household file
SHS_COLS = ["SEQNO", "CNTYFIPS", "PHD", "INCCAT", "ADULTS", "CHILDREN", "AGEGRP", "GENHLTH", "MEDEQUIP", "HH_WT"]


def shs_rows(x):
    B = x["B"]
    wb = S.final_weights(B["base"], B["county"], B["band"], x["pub"])
    rng = np.random.default_rng([S.SEED, 22])
    order = rng.permutation(len(B["base"]))
    ids = numbers(2, len(order), 1, 99999)
    rows = []
    for k, i in enumerate(order):
        c = B["county"][i]
        rows.append([2025000000 + int(ids[k]), int(W.FIPS[c]), W.PHD[c], int(B["cat"][i]), int(B["adults"][i]),
                     int(B["children"][i]), int(B["age"][i]), int(B["genhlth"][i]), int(B["answer"][i]), weight(wb[i])])
    return rows


# ---------------------------------------------------------------- the screening pilot
PILOT_COLS = ["SCREEN_ID", "ADDRESS_KEY", "DIVISION", "MODE", "RECEIVED", "STATUS", "HH_INCOME", "ELEC_12M",
              "OTHER_FUELS_12M", "MED_EQUIPMENT"]


def pilot_rows(x):
    rows = []
    ids = numbers(3, len(x["screens"]), 100001, 199999)
    for k, r in enumerate(x["screens"]):
        complete = r["status"] in ("Complete", "Partial")
        rows.append(["S25-%06d" % ids[k], r["key"], r["division"], r["mode"], r["received"].isoformat(),
                     r["status"], int(r["income"]), int(r["elec"]) if complete else "",
                     int(r["other"]) if complete else "", r["equip"] or ""])
    return rows


# ---------------------------------------------------------------- the division register
DIV_COLS = ["COUNTY_FIPS", "COUNTY", "SERVICE_DIVISION"]


def division_rows():
    return [[W.FIPS[c], c, W.DIVISION_OF[c]] for c in W.COUNTIES if c in W.DIVISION_OF]


# ---------------------------------------------------------------- the medical baseline enrolment extract
ENROL_COLS = ["ENROL_ID", "ACCOUNT_NO", "PRIOR_ACCOUNT_NO", "PREMISE_ID", "SERVICE_COUNTY", "DIVISION", "START_DATE",
              "END_DATE", "END_REASON", "QUALIFYING_EQUIPMENT", "LIFE_SUPPORT"]


def enrol_rows(x):
    spells = []
    for hh in x["households"]:
        for sp in hh["spells"]:
            spells.append((sp["start"], sp["account"], hh, sp))
    spells.sort(key=lambda t: (t[0], t[1]))
    rows = []
    ids = numbers(4, len(spells), 1000001, 1999999)
    for k, (_, _, hh, sp) in enumerate(spells):
        rows.append(["MB%07d" % ids[k], sp["account"], sp["prior"], sp["premise"], sp["county"],
                     sp["division_stored"], sp["start"].isoformat(), sp["end"].isoformat() if sp["end"] else "",
                     sp["reason"], hh["equipment"], "Y" if hh["life_support"] else "N"])
    return rows


# ---------------------------------------------------------------- the battery work orders
ORDER_COLS = ["WORK_ORDER", "ACCOUNT_NO", "PREMISE_ID", "SERVICE_COUNTY", "SCHEDULED", "STATUS", "COMPLETED_UTC",
              "UNIT_SERIAL", "REMOVED_SERIAL", "CREW"]


def order_rows(x):
    rng = np.random.default_rng([S.SEED, 23])
    crews = {d: ["%s-%02d" % ("".join(w[0] for w in d.split()), k) for k in range(1, 4)] for d in W.DIV_NAMES}
    rows = []
    ids = numbers(15, len(x["orders"]), 410001, 499999)
    for k, o in enumerate(x["orders"]):
        d = W.DIVISION_OF[o["sp"]["county"]]
        at = o["at"].strftime("%Y-%m-%dT%H:%M:%SZ") if o["at"] else ""
        rows.append(["WO-%06d" % ids[k], o["sp"]["account"], o["sp"]["premise"], o["sp"]["county"],
                     o["day"].isoformat(), o["status"], at, o["unit"], o["removed"],
                     crews[d][int(rng.integers(0, 3))]])
    return rows


# ---------------------------------------------------------------- the two neighbouring state files
HVI_COLS = ["COUNTY_FIPS", "COUNTY", "HVI_SCORE", "PCT_AGE_65_PLUS", "PCT_65_PLUS_LIVING_ALONE", "PCT_NO_CENTRAL_AC",
            "PCT_IMPERVIOUS", "HVI_RANK"]


def hvi_rows(x):
    return [[W.FIPS[h["county"]], h["county"], "%.1f" % h["score"], "%.1f" % h["old"], "%.1f" % h["alone"],
             "%.1f" % h["noac"], "%.1f" % h["imperv"], h["rank"]] for h in x["hvi"]]


APPS_COLS = ["COUNTY_FIPS", "COUNTY", "MONTH", "APPLICATIONS_RECEIVED", "APPLICATIONS_APPROVED", "CRISIS_APPLICATIONS",
             "AVERAGE_BENEFIT_USD"]


def apps_rows(x):
    return [[W.FIPS[a["county"]], a["county"], a["month"], a["received"], a["approved"], a["crisis"], a["benefit"]]
            for a in x["apps"]]


# ---------------------------------------------------------------- the demographer's household estimates
def _sheet_header(ws, title, sub, row=1):
    from openpyxl.styles import Font
    ws.cell(row=row, column=1, value=title).font = Font(name="Calibri", size=13, bold=True)
    ws.cell(row=row + 1, column=1, value=sub).font = Font(name="Calibri", size=9, italic=True, color="595959")


def estimates_workbook(path, x):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    pub = x["pub"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Households by income"
    _sheet_header(ws, "Households by household income, counties and planning regions: July 1, 2025",
                  "State Demographer's Office, county household estimates, vintage 2025. Released March 18, 2026.")
    head = ["Area", "FIPS", "Total households", "Under $25,000", "$25,000 to $49,999", "$50,000 to $74,999",
            "$75,000 and over"]
    thin = Side(style="thin", color="A6A6A6")
    for j, h in enumerate(head, start=1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(name="Calibri", size=10, bold=True)
        c.fill = PatternFill("solid", fgColor="DCE6F0")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = Border(bottom=thin)
    bands = lambda c: [int(pub[(c, b)]) for b in range(4)]
    regions = sorted(set(W.REGION.values()))
    state = [sum(bands(c)[b] for c in W.COUNTIES) for b in range(4)]
    row = [5]

    def put(label, fips, vals, bold=False, indent=0):
        cells = [label, fips, sum(vals)] + vals
        for j, v in enumerate(cells, start=1):
            c = ws.cell(row=row[0], column=j, value=v)
            c.font = Font(name="Calibri", size=10, bold=bold)
            if j >= 3:
                c.number_format = "#,##0"
            if j == 1 and indent:
                c.alignment = Alignment(indent=indent)
        row[0] += 1
    put("State", "000", state, bold=True)
    for g in regions:
        cs = [c for c in W.COUNTIES if W.REGION[c] == g]
        put("Planning Region %d" % g, None, [sum(bands(c)[b] for c in cs) for b in range(4)], bold=True)
        for c in cs:
            put("%s County" % c, W.FIPS[c], bands(c), indent=1)
    widths = [26, 7, 14, 13, 14, 14, 13]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=4, column=j).column_letter].width = w
    ws.freeze_panes = "A5"
    r = row[0] + 1
    for line in ("Households are occupied housing units. Household income is money income in the previous 12 months.",
                 "Estimates are rounded to the nearest ten; region and state rows are the sums of their county rows.",
                 "Source: State Demographer's Office, county household estimates program, vintage 2025."):
        ws.cell(row=r, column=1, value=line).font = Font(name="Calibri", size=9, italic=True, color="595959")
        r += 1
    wb.save(path)
    return len(W.COUNTIES)
