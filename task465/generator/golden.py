"""The goldens, computed from the shipped pack only (bundle.py) with the fusion in methods.py: the staffing workbook
with its chart and the conformed table of eligible households; generator/figures.json carries every figure the
write-up states."""
import io
import json
import math
import os

import bundle as B
import methods as M
import world as W


def r100(v):
    return int(math.floor(v / 100.0 + 0.5) * 100)


def compute(target):
    N, p, q, tabs, tabs_div = B.conformed(target)
    R, info = M.routes(N, p, q, tabs, tabs_div)
    g = R["golden"]
    gd = M.division_totals(g)
    staff = {d: M.caseworkers(gd[d]) for d in W.DIV_NAMES}
    rivals = {}
    for m in ("independence", "pilot_shares", "energy_x_pilot", "health_x_pilot", "lift"):
        dt_ = M.division_totals(R[m])
        rivals[m] = dict(caseworkers={d: M.caseworkers(dt_[d]) for d in W.DIV_NAMES},
                         total=sum(M.caseworkers(dt_[d]) for d in W.DIV_NAMES), eligible=r100(sum(dt_.values())))
    pooled = M.independence_by_division(N, p, q)
    rivals["independence_pooled"] = dict(total=sum(M.caseworkers(v) for v in pooled.values()))
    rates = {}
    for b in W.ELIGIBLE_BANDS:
        t = tabs[b]
        rates[str(b)] = dict(equip_high=round(100 * t[0] / (t[0] + t[1]), 1), equip_rest=round(100 * t[2] / (t[2] + t[3]), 1),
                             screens=sum(t))
    F = {
        "call": {"total": sum(staff.values()), "by_division": staff},
        "eligible": {d: r100(gd[d]) for d in W.DIV_NAMES},
        "eligible_sum": sum(r100(gd[d]) for d in W.DIV_NAMES),
        "eligible_raw": {d: round(gd[d], 2) for d in W.DIV_NAMES},
        "cells": {"%s|%d" % k: r100(v) for k, v in g.items()},
        "cells_raw": {"%s|%d" % k: round(v, 2) for k, v in g.items()},
        "households": {"%s|%d" % k: int(v) for k, v in N.items()},
        "burden_share": {"%s|%d" % k: round(v, 4) for k, v in p.items()},
        "equipment_share": {"%s|%d" % k: round(v, 4) for k, v in q.items()},
        "psi": {str(b): round(info["psi_band"][b], 2) for b in W.ELIGIBLE_BANDS},
        "psi_mh": round(info["psi_mh"], 2),
        "pilot_rates": rates,
        "rivals": rivals,
        "baseline": B.baseline_households(target),
        "battery": B.battery_households(target),
    }
    F["stated"] = {"caseload": W.CASELOAD, "band_edges": [25000, 49999, 50000], "burden_line_pct": 6,
                   "data_year": W.DATA_YEAR, "plan_year": W.PLAN_YEAR, "count_date": "2026-10-31"}
    F["baseline_total"] = sum(F["baseline"].values())
    F["battery_total"] = sum(F["battery"].values())
    return F


def chart_png(F):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    divs = sorted(W.DIV_NAMES, key=lambda d: (F["eligible"][d], d))
    vals = [F["eligible"][d] for d in divs]
    fig, ax = plt.subplots(figsize=(8.2, 4.6), dpi=130)
    bars = ax.barh(divs, vals, color="#2f6f73", height=0.62, zorder=2)
    for b, d in zip(bars, divs):
        n = F["call"]["by_division"][d]
        ax.text(b.get_width() + 30, b.get_y() + b.get_height() / 2, "%d caseworkers" % n, va="center", fontsize=9,
                color="#1b1d1f")
    ax.set_xlim(0, max(vals) * 1.22)
    ax.set_xlabel("Eligible households, 2025 (nearest hundred)", fontsize=9, color="#3b3f44")
    ax.set_title("HEMA eligible households and 2027 caseworkers by service division", fontsize=11, loc="left",
                 color="#1b1d1f")
    ax.grid(axis="x", color="#d9dde1", linewidth=0.8, zorder=0)
    ax.tick_params(axis="both", labelsize=9, colors="#3b3f44")
    ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: "{:,}".format(int(v))))
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def write_workbook(path, F):
    import openpyxl
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Caseworkers 2027"
    ws["A1"] = "HEMA caseworkers for 2027 by service division"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="1F3A5F")
    ws["A2"] = ("Eligible households are 2025 counts from the conformed table, rounded to the nearest hundred. Caseworkers "
                "follow the plan's standard: one for every 400 eligible households in a division, or part of that number.")
    ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="595959")
    head = ["Service division", "Eligible households, 2025", "Caseworkers, 2027",
            "Households on the medical baseline rate, 31 October 2026",
            "Households given a backup battery, January to October 2026"]
    thin = Side(style="thin", color="A6A6A6")
    for j, h in enumerate(head, start=1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(name="Calibri", size=10, bold=True)
        c.fill = PatternFill("solid", fgColor="DDE8EE")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(bottom=thin)
    r = 5
    for d in W.DIV_NAMES:
        vals = [d, F["eligible"][d], F["call"]["by_division"][d], F["baseline"][d], F["battery"][d]]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = Font(name="Calibri", size=10)
            if j > 1:
                c.number_format = "#,##0"
        r += 1
    tot = ["Total", F["eligible_sum"], F["call"]["total"], F["baseline_total"], F["battery_total"]]
    for j, v in enumerate(tot, start=1):
        c = ws.cell(row=r, column=j, value=v)
        c.font = Font(name="Calibri", size=10, bold=True)
        c.border = Border(top=thin)
        if j > 1:
            c.number_format = "#,##0"
    for col, w in zip("ABCDE", (18, 14, 12, 22, 24)):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[4].height = 46
    img = Image(io.BytesIO(chart_png(F)))
    img.width, img.height = 640, 359
    ws.add_image(img, "A17")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws2 = wb.create_sheet("Notes")
    notes = [
        ("Eligible households", "Households with incomes under $50,000 whose annual home energy costs exceed 6 per cent "
         "of income and in which someone relies on electrically powered medical equipment at home (plan, section 2), "
         "2025, by division and income band in hema_eligible_households_2025.csv."),
        ("Conformed table", "Each division and band's households are the State Demographer's July 1, 2025 estimates; "
         "its high-burden share is the Household Energy Affordability Survey 2025 estimate (WT_FINAL) and its equipment "
         "share the State Health Survey 2025 estimate (HH_WT, don't know and refused excluded). The two traits are "
         "joined with the screening pilot's odds ratio inside each band (%.2f under $25,000, %.2f from $25,000 to "
         "$49,999), the pilot's band table raked to each cell's two shares." % (F["psi"]["0"], F["psi"]["1"])),
        ("Caseworkers", "From each division's unrounded eligible households."),
        ("Medical baseline households", "Households with a medical baseline spell open on October 31, 2026, accounts "
         "joined through transfers on a move, each in the division of its premise's county."),
        ("Backup battery households", "Households with a completed placement on a Central Time date from January 1 to "
         "October 31, 2026, accounts joined through transfers, each in the division of its premise's county."),
    ]
    ws2["A1"] = "Notes"
    ws2["A1"].font = Font(name="Calibri", size=12, bold=True)
    for i, (k, v) in enumerate(notes, start=3):
        ws2.cell(row=i, column=1, value=k).font = Font(name="Calibri", size=10, bold=True)
        c = ws2.cell(row=i, column=2, value=v)
        c.font = Font(name="Calibri", size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws2.column_dimensions["A"].width = 26
    ws2.column_dimensions["B"].width = 100
    wb.save(path)


def write_table(path, F):
    import csv
    with open(path, "w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh, lineterminator="\r\n")
        wr.writerow(["SERVICE_DIVISION", "INCOME_BAND", "HOUSEHOLDS_2025", "ELIGIBLE_HOUSEHOLDS_2025"])
        for d in W.DIV_NAMES:
            for b in W.ELIGIBLE_BANDS:
                k = "%s|%d" % (d, b)
                wr.writerow([d, W.BANDS[b], F["households"][k], F["cells"][k]])


def main(target, gdir, fig_path):
    os.makedirs(gdir, exist_ok=True)
    F = compute(target)
    write_workbook(os.path.join(gdir, W.GOLDEN["workbook"]), F)
    write_table(os.path.join(gdir, W.GOLDEN["table"]), F)
    with open(fig_path, "w", encoding="utf-8") as fh:
        json.dump(F, fh, indent=1, sort_keys=True)
        fh.write("\n")
    return F
