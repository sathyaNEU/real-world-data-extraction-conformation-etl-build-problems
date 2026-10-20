"""The pack's data files, written from the world: the IRS ZIP-code extracts, the relationship files, the job export,
its status history, the grants ledger and the lookups."""
import csv
import datetime as dt

import numpy as np

import world as W

FILES = {
    "irs23": "irs_soi_zip_code_data_ty2023_extract.csv",
    "irs22": "irs_soi_zip_code_data_ty2022_extract.csv",
    "rel": "zip_county_relationship_2023q4.xlsx",
    "rel_rev": "county_zip_relationship_2023q4.xlsx",
    "dor": "dor_individual_income_tax_statistics_ty2023.pdf",
    "policy": "cvcog_weatherization_allocation_policy_rev4.pdf",
    "members": "cvcog_member_counties.csv",
    "agencies": "wx_subgrantee_agencies.csv",
    "jobs": "wx_job_export_2026-09-30.csv",
    "history": "wx_job_status_history_2026-09-30.csv",
    "ledger": "cvcog_grants_ledger_2024-07_2026-09.csv",
    "closeout": "cvcog_wap_py2025_closeout_summary.pdf",
    "dictionary": "data_dictionary.txt",
    "manifest": "export_manifest.csv",
}
USERS = {"THCA": ["THCA-intake1", "THCA-dlong", "THCA-crew2"], "NCCS": ["NCCS-intake", "NCCS-mwade", "NCCS-crewlead"],
         "UCAP": ["UCAP-front", "UCAP-jpark", "UCAP-qc"], "EHCA": ["EHCA-intake", "EHCA-rmoss", "EHCA-field"],
         "SCCA": ["SCCA-intake", "SCCA-tbell", "SCCA-crew1"]}


def write_csv(path, cols, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\r\n")
        w.writerow(cols)
        for r in rows:
            w.writerow(r)
    return len(rows)


# ---------------------------------------------------------------------------------------------- IRS extracts
IRS23_COLS = ["STATEFIPS", "ZIPCODE", "AGI_STUB", "N1", "A00100", "N59660", "A59660"]
IRS22_COLS = ["statefips", "zipcode", "agi_stub", "n1", "a00100", "n59660", "a59660"]


def irs_rows(pub):
    zips = sorted({z for z, _ in pub if z not in ("00000", "99999")})
    out = []
    for z in ["00000"] + zips + ["99999"]:
        for s in range(1, 7):
            n1, agi, ec, ea = pub[(z, s)]
            out.append([W.STATE_FIPS, z, s, n1, agi, ec, ea])
    return out


# ---------------------------------------------------------------------------------------------- relationship files
REL_COLS = ["ZIP", "COUNTY", "RES_RATIO", "BUS_RATIO", "OTH_RATIO", "TOT_RATIO"]
REV_COLS = ["COUNTY", "ZIP", "RES_RATIO", "BUS_RATIO", "OTH_RATIO", "TOT_RATIO"]


def rel_rows(rel):
    return [[r["zip"], r["fips"], r["res"], r["bus"], r["oth"], r["tot"]] for r in rel]


def rev_rows(rel):
    """The same addresses the other way: each county's addresses by ZIP code."""
    by = {}
    for r in rel:
        by.setdefault(r["fips"], []).append(r)
    out = []
    for f in sorted(by):
        rs = by[f]
        tot = [sum(r["addr"][i] for r in rs) for i in range(3)]
        for r in sorted(rs, key=lambda r: r["zip"]):
            a = r["addr"]
            ratio = [round(a[i] / tot[i], 6) if tot[i] else 0.0 for i in range(3)]
            out.append([f, r["zip"], ratio[0], ratio[1], ratio[2], round(sum(a) / sum(tot), 6) if sum(tot) else 0.0])
    return out


# ---------------------------------------------------------------------------------------------- the job system
JOB_COLS = ["JOB_NO", "AGENCY", "COUNTY", "TRACT", "BUILDING_TYPE", "UNITS", "FUNDING", "HEATING_FUEL", "STATUS",
            "OPENED", "AUDIT_DATE", "START_DATE", "COMPLETED_DATE"]
HIST_COLS = ["JOB_NO", "STATUS", "CHANGED_AT", "CHANGED_BY", "NOTE"]


def job_rows(x):
    out = []
    for j in sorted(x["jobs"], key=lambda j: j["no"]):
        st, last = W.status_now(j)
        county = "Tallis" if j["agency"] == "THCA" else W.NAME[j["tract"][:5]]
        out.append([j["no"], j["agency"], county, j["tract"], j["btype"], j["units"], j["funding"],
                    j["fuel"], st, j["created"].isoformat(), j["audit"].isoformat() if j["audit"] else "",
                    j["start"].isoformat() if j["start"] else "", last.isoformat() if last else ""])
    return out


def stamp(rng, d):
    return dt.datetime.combine(d, dt.time(int(rng.integers(7, 18)), int(rng.integers(0, 60))))


NOTES = {
    "AUDIT SCHEDULED": ["Phone", "Walk-in", "Referral: LIHEAP", "Referral: utility", "Phone", "Mail-in application"],
    "DEFERRED": ["Roof leak, referred to home repair", "Knob and tube wiring, referred out", "Client asked to defer",
                 "Moisture in crawlspace", "Vermiculite in attic", "Structural, referred to home repair"],
    "CANCELLED": ["Client moved", "Over income at recertification", "Client withdrew", "Landlord declined",
                  "No contact after three attempts"],
    "REOPENED": ["Customer call: draft at front door", "Customer call: CO alarm sounding",
                 "Failed customer check: attic hatch", "Customer call: furnace short cycling",
                 "Customer call: bath fan not venting", "Customer call: window will not lock",
                 "Inspector return: duct joint open", "Customer call: water heater pilot out"],
    "SCHEDULED": ["Crew 1", "Crew 2", "Crew 3", "Sub: furnace", "Crew 2", "Crew 4"],
}
ORDER = ["OPENED", "AUDIT SCHEDULED", "AUDITED", "DEFERRED", "CANCELLED", "APPROVED", "SCHEDULED", "IN PROGRESS",
         "COMPLETED", "REOPENED", "MOVED"]


def history_rows(x):
    """Every status change the job system holds; the jobs moved from the old system carry their old history under
    the new number, and the move itself is a change whose note gives the old number."""
    rng = np.random.default_rng(W.SEED + 920)
    out = []
    for j in sorted(x["jobs"], key=lambda j: j["no"]):
        u = USERS[j["agency"]]
        ev = [("OPENED", j["created"], u[0]), ("AUDIT SCHEDULED", W.day(j["created"], rng.integers(0, 2)), u[0])]
        if j["audit"]:
            ev.append(("AUDITED", j["audit"], u[1]))
        if j["deferred"]:
            ev.append(("DEFERRED", j["deferred"], u[1]))
        elif j["cancelled"]:
            ev.append(("CANCELLED", j["cancelled"], u[0]))
        elif j["start"]:
            ev.append(("APPROVED", W.day(j["audit"], rng.integers(1, 6)), u[1]))
            ev.append(("SCHEDULED", W.day(j["start"], -int(rng.integers(3, 11))), u[2]))
            ev.append(("IN PROGRESS", j["start"], u[2]))
            for k, c in enumerate(j["completions"]):
                ev.append(("COMPLETED", c, u[2]))
                if k < len(j["reopens"]):
                    ev.append(("REOPENED", j["reopens"][k], u[1]))
                    ev.append(("SCHEDULED", W.day(j["reopens"][k], rng.integers(1, 4)), u[2]))
        ev = [e for e in ev if e[1] <= W.PULL]
        last = None
        rows = []
        for st, d, who in sorted(ev, key=lambda e: (e[1], ORDER.index(e[0]))):
            t = stamp(rng, d)
            if last and t <= last:
                t = last + dt.timedelta(minutes=int(rng.integers(5, 90)))
            last = t
            note = ""
            if st in NOTES and (st != "SCHEDULED" or rng.random() < 0.45) and (st != "AUDIT SCHEDULED"
                                                                                or rng.random() < 0.6):
                note = NOTES[st][int(rng.integers(len(NOTES[st])))]
            rows.append([j["no"], st, t, who, note])
        if j["legacy"]:
            t = dt.datetime.combine(W.MOVE, dt.time(5, 0)) + dt.timedelta(seconds=int(j["no"] - W.NEW_FIRST) * 7)
            rows.append([j["no"], "MOVED", t, "conversion", "Converted; previous job no. %d" % j["legacy"]])
            rows.sort(key=lambda r: r[2])
        for r in rows:
            out.append([r[0], r[1], r[2].strftime("%Y-%m-%d %H:%M"), r[3], r[4]])
    return out


# ---------------------------------------------------------------------------------------------- the ledger
LEDGER_COLS = ["LINE_ID", "POSTED_DATE", "AGENCY", "INVOICE_NO", "JOB_NO", "ACCOUNT", "FUND", "SERVICE_DATE",
               "DESCRIPTION", "AMOUNT"]


def ledger_rows(x):
    rng = np.random.default_rng(W.SEED + 930)
    out = []
    lid = 4410217
    for ln in x["ledger"]:
        lid += int(rng.integers(1, 4))
        acct = "LABOR" if ln["acct"] == "CALLBACK" else ln["acct"]
        texts = W.ACCT_TEXT.get(ln["acct"])
        desc = texts[int(rng.integers(len(texts)))] if texts else {"ADMIN": "Administration", "VEHICLE": "Vehicle costs",
                                                                     "TRAINING": "Crew training",
                                                                     "INSURANCE": "Liability insurance"}[ln["acct"]]
        out.append([lid, ln["posted"].isoformat(), ln["agency"], ln["invoice"], ln["job_no"] or "", acct, ln["fund"],
                    ln["service"].isoformat(), desc, "%.2f" % ln["amount"]])
    return out


# ---------------------------------------------------------------------------------------------- lookups
MEMBER_COLS = ["COUNTY_FIPS", "COUNTY", "AGENCY"]
AGENCY_COLS = ["AGENCY", "AGENCY_NAME", "COUNTIES_SERVED", "CONTRACT_NO", "OFFICE_COUNTY"]


def member_rows():
    return [[f, W.NAME[f], W.AGENCY_OF[f]] for f in W.FIPS]


def agency_rows():
    out = []
    for k, (a, (nm, fs)) in enumerate(W.AGENCIES.items()):
        out.append([a, nm, "; ".join(W.NAME[f] for f in fs), "WX-26-%02d" % (k + 1), W.NAME[fs[0]]])
    return out


def write_xlsx(path, sheet, cols, rows, ratio_cols=()):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(list(r))
    for j in ratio_cols:
        for row in ws.iter_rows(min_row=2, min_col=j + 1, max_col=j + 1):
            for cell in row:
                cell.number_format = "0.000000"
    for k, col in enumerate(cols):
        ws.column_dimensions[openpyxl.utils.get_column_letter(k + 1)].width = 12 if k < 2 else 13
    ws.freeze_panes = "A2"
    wb.save(path)
    return len(rows)
