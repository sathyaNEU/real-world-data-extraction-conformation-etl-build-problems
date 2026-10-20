"""Independent verifier: recomputes every graded figure from the shipped pack with its own code (pandas for the
tables, pypdf for the Department's report, its own allocation loop), never importing bundle.py or world.py, and
compares with generator/figures.json and the golden files.

    python3 task475/generator/verify.py [target] [golden]
"""
import json
import os
import re
import sys

import numpy as np
import openpyxl
import pandas as pd
import pypdf

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
LABELS = ["Under $10,000", "$10,000 under $20,000", "$20,000 under $25,000", "$25,000 under $50,000",
          "$50,000 under $75,000", "$75,000 or more"]


def dor_table(path):
    txt = "\n".join(pg.extract_text() for pg in pypdf.PdfReader(path).pages)
    lines = [ln.strip() for ln in txt[txt.index("Table 2."):].splitlines()]
    out = {}
    for i, ln in enumerate(lines):
        if ln in LABELS:
            v = [int(x.replace(",", "")) for x in lines[i + 1:i + 9]]
            out[LABELS.index(ln) + 1] = v
    assert len(out) == 6
    return out


def families(target):
    irs = pd.read_csv(os.path.join(target, "irs_soi_zip_code_data_ty2023_extract.csv"), dtype={"ZIPCODE": str})
    irs = irs[~irs["ZIPCODE"].isin(["00000", "99999"])]
    rel = pd.read_excel(os.path.join(target, "zip_county_relationship_2023q4.xlsx"), dtype={"ZIP": str, "COUNTY": str})
    res_tot = rel.groupby("ZIP")["RES_RATIO"].transform("sum")
    rel["W"] = np.where(res_tot > 0, rel["RES_RATIO"], rel["TOT_RATIO"])
    t2 = dor_table(os.path.join(target, "dor_individual_income_tax_statistics_ty2023.pdf"))
    first = [t2[c] for c in (1, 2, 3)]
    n0 = sum(r[0] for r in first)
    s0 = sum(r[1] for r in first) * 1000.0
    n1 = sum(r[2] + r[4] + r[6] for r in first)
    s1 = sum(r[3] + r[5] + r[7] for r in first) * 1000.0
    a0, a1 = s0 / n0, s1 / n1
    assert all(t2[c][0] == 0 for c in (4, 5, 6))
    irs = irs.copy()
    irs["FAM"] = np.where(irs["AGI_STUB"] == 1, (irs["A59660"] * 1000.0 - a0 * irs["N59660"]) / (a1 - a0),
                          irs["N59660"])
    z = irs.groupby("ZIPCODE")["FAM"].sum().rename("FAM").reset_index()
    m = rel.merge(z, left_on="ZIP", right_on="ZIPCODE", how="inner")
    m["V"] = m["W"] * m["FAM"]
    members = pd.read_csv(os.path.join(target, "cvcog_member_counties.csv"), dtype={"COUNTY_FIPS": str})
    fam = m.groupby("COUNTY")["V"].sum()
    return {r.COUNTY: float(fam[r.COUNTY_FIPS]) for r in members.itertuples()}


def allocate(fam, fund=20_000_000.0, cap=0.12):
    left = dict(fam)
    out = {}
    pot = fund
    while True:
        tot = sum(left.values())
        trial = {k: pot * v / tot for k, v in left.items()}
        over = [k for k, v in trial.items() if v > cap * fund]
        if not over:
            out.update(trial)
            return out
        for k in over:
            out[k] = cap * fund
            pot -= cap * fund
            del left[k]


def asks(target):
    jobs = pd.read_csv(os.path.join(target, "wx_job_export_2026-09-30.csv"), dtype={"TRACT": str})
    hist = pd.read_csv(os.path.join(target, "wx_job_status_history_2026-09-30.csv"), keep_default_na=False)
    led = pd.read_csv(os.path.join(target, "cvcog_grants_ledger_2024-07_2026-09.csv"), keep_default_na=False)
    members = pd.read_csv(os.path.join(target, "cvcog_member_counties.csv"), dtype={"COUNTY_FIPS": str})
    name = dict(zip(members["COUNTY_FIPS"], members["COUNTY"]))
    jobs["CTY"] = jobs["TRACT"].str[:5].map(name)
    prog = jobs["FUNDING"].str.contains("DOE-WAP|STATE-WX")
    first = hist[hist["STATUS"] == "COMPLETED"].groupby("JOB_NO")["CHANGED_AT"].min().str[:10]
    j = jobs[prog].copy()
    j["FIRST"] = j["JOB_NO"].map(first)
    done = j[(j["FIRST"] >= "2025-01-01") & (j["FIRST"] <= "2025-12-31")]
    a = done.groupby("CTY")["UNITS"].sum()
    moved = hist[hist["STATUS"] == "MOVED"].copy()
    moved["OLD"] = moved["NOTE"].str.extract(r"previous job no\. (\d+)")[0].astype(int)
    old_to_new = dict(zip(moved["OLD"], moved["JOB_NO"]))
    l25 = led[(led["POSTED_DATE"] >= "2025-01-01") & (led["POSTED_DATE"] <= "2025-12-31") & (led["JOB_NO"] != "")]
    assert (l25["POSTED_DATE"] < "2026-01-12").all()
    l25 = l25.copy()
    l25["JOB"] = l25["JOB_NO"].astype(int).map(old_to_new)
    assert l25["JOB"].notna().all()
    l25["CTY"] = l25["JOB"].map(dict(zip(jobs["JOB_NO"], jobs["CTY"])))
    b = l25.groupby("CTY")["AMOUNT"].sum()
    return {k: int(a.get(k, 0)) for k in name.values()}, {k: float(b.get(k, 0.0)) for k in name.values()}


def main(target=None, gdir=None):
    target = target or os.path.join(TASK, "target")
    gdir = gdir or os.path.join(TASK, "golden")
    fig = json.load(open(os.path.join(HERE, "figures.json")))
    comp = fig["components"]
    fam = families(target)
    al = allocate(fam)
    tot = sum(fam.values())
    bad = []
    for k in fam:
        if round(fam[k], 1) != comp["families"][k]:
            bad.append(("families", k, fam[k]))
        if round(100.0 * fam[k] / tot, 2) != comp["shares_pct"][k]:
            bad.append(("share", k, 100.0 * fam[k] / tot))
        if int(round(al[k] / 100.0)) * 100 != comp["allocations"][k]:
            bad.append(("allocation", k, al[k]))
    capped = [k for k, v in al.items() if abs(v - 2_400_000) < 1e-6]
    if (capped or ["none"]) != [fig["call"]["capped"]]:
        bad.append(("capped", capped))
    a, b = asks(target)
    for k in a:
        if a[k] != fig["a_dwellings"][k]:
            bad.append(("dwellings", k, a[k]))
        if int(round(b[k])) != fig["b_spending"][k]:
            bad.append(("spending", k, b[k]))
    if os.path.isdir(gdir):
        wb = openpyxl.load_workbook(os.path.join(gdir, "corran_valley_2027_weatherization_allocation.xlsx"))
        seen = 0
        for row in wb["Allocation"].iter_rows(values_only=True):
            if row[0] in comp["allocations"]:
                seen += 1
                if (row[3], row[4], row[5]) != (comp["families"][row[0]], comp["shares_pct"][row[0]],
                                                comp["allocations"][row[0]]):
                    bad.append(("workbook", row[0]))
        for row in wb["Program year 2025"].iter_rows(min_row=2, max_row=15, values_only=True):
            if (row[1], row[2]) != (fig["a_dwellings"][row[0]], fig["b_spending"][row[0]]):
                bad.append(("workbook asks", row[0]))
        csvt = pd.read_csv(os.path.join(gdir, "corran_valley_ty2023_county_credit_families.csv"))
        for r in csvt.itertuples():
            if round(r.EITC_RETURNS_WITH_QUALIFYING_CHILDREN, 1) != comp["families"][r.COUNTY]:
                bad.append(("csv", r.COUNTY))
        if seen != 14:
            bad.append(("workbook rows", seen))
    n = 14 * 5 + 1
    if bad:
        print("verify: %d of %d figures differ: %s" % (len(bad), n, bad[:10]))
        raise SystemExit(1)
    print("verify: all %d graded figures recompute independently (families, shares, allocations, the cap, dwellings, "
          "spending) and the goldens carry them" % n)


if __name__ == "__main__":
    main(*(sys.argv[1:3]))
