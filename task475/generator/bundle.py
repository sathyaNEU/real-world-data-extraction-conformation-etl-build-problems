"""Read the shipped pack and build the council's figures from it: each member county's tax-year 2023 earned income
credit returns with qualifying children, conformed from the IRS ZIP-code statistics through the residential ZIP to
county file and the Department of Revenue's statewide table; the capped allocation of the 2027 fund; and the two asks
over program year 2025 (dwellings weatherized, spending on dwellings). Only the pack's files are read. The routes
replay the easy path, the natural mishandlings and the wrong readings; checks.py and ceiling.py run them."""
import csv
import datetime as dt
import os
import re

import numpy as np
import openpyxl

FILES = dict(irs23="irs_soi_zip_code_data_ty2023_extract.csv", irs22="irs_soi_zip_code_data_ty2022_extract.csv",
             rel="zip_county_relationship_2023q4.xlsx", rel_rev="county_zip_relationship_2023q4.xlsx",
             dor="dor_individual_income_tax_statistics_ty2023.pdf",
             policy="cvcog_weatherization_allocation_policy_rev4.pdf", members="cvcog_member_counties.csv",
             agencies="wx_subgrantee_agencies.csv", jobs="wx_job_export_2026-09-30.csv",
             history="wx_job_status_history_2026-09-30.csv", ledger="cvcog_grants_ledger_2024-07_2026-09.csv",
             closeout="cvcog_wap_py2025_closeout_summary.pdf")
# the policy: the fund (the 2027 budget line), the cap (5.2), rounding (5.3); the record of the last closed year (6.1)
FUND = 20_000_000
CAP = 0.12
PY0, PY1 = dt.date(2025, 1, 1), dt.date(2025, 12, 31)
MOVE = dt.date(2026, 1, 12)
PROGRAM = ("DOE-WAP", "STATE-WX")
DOR_LABELS = ["Under $10,000", "$10,000 under $20,000", "$20,000 under $25,000", "$25,000 under $50,000",
              "$50,000 under $75,000", "$75,000 or more"]
IRS_CLASS1_DOR = (1, 2, 3)          # the IRS's first class, $1 under $25,000, is the Department's classes 1 to 3

_CACHE = {}


def p(target, key):
    return os.path.join(target, FILES[key])


def rows_csv(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def cached(key, fn):
    if key not in _CACHE:
        _CACHE[key] = fn()
    return _CACHE[key]


# ---------------------------------------------------------------------------------------------- readers
def members(target):
    return [(r["COUNTY_FIPS"], r["COUNTY"], r["AGENCY"]) for r in rows_csv(p(target, "members"))]


def member_fips(target):
    return [m[0] for m in members(target)]


def names(target):
    return {f: n for f, n, _ in members(target)}


def irs(target, key="irs23"):
    """{(zip, class): (returns, AGI $000, credit returns, credit $000)} in either year's layout."""
    def build():
        out = {}
        for r in rows_csv(p(target, key)):
            r = {k.upper(): v for k, v in r.items()}
            out[(r["ZIPCODE"].zfill(5), int(r["AGI_STUB"]))] = (int(r["N1"]), int(r["A00100"]), int(r["N59660"]),
                                                               int(r["A59660"]))
        return out
    return cached(("irs", target, key), build)


def relationship(target, key="rel"):
    """The relationship file's rows: ZIP code to county (rel) or county to ZIP code (rel_rev), with its four ratios."""
    def build():
        wb = openpyxl.load_workbook(p(target, key), read_only=True)
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        head = list(next(it))
        out = []
        for row in it:
            r = dict(zip(head, row))
            out.append(dict(zip=str(r["ZIP"]).zfill(5), fips=str(r["COUNTY"]).zfill(5), res=float(r["RES_RATIO"]),
                            bus=float(r["BUS_RATIO"]), oth=float(r["OTH_RATIO"]), tot=float(r["TOT_RATIO"])))
        wb.close()
        return out
    return cached(("rel", target, key), build)


def pdf_text(path):
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        return "\n".join(pg.extract_text() or "" for pg in pdf.pages)


def dor(target):
    """The Department's Table 2: {(its class 1 to 6, children 0 to 3): (credit returns, credit $000)}."""
    def build():
        txt = pdf_text(p(target, "dor"))
        out = {}
        num = r"([\d,]+)"
        for line in txt.splitlines():
            for c, lab in enumerate(DOR_LABELS, start=1):
                m = re.match(re.escape(lab) + r"\s+" + r"\s+".join([num] * 8) + r"\s*$", line.strip())
                if m:
                    v = [int(g.replace(",", "")) for g in m.groups()]
                    for k in range(4):
                        out[(c, k)] = (v[2 * k], v[2 * k + 1])
        assert len(out) == 24, "the Department's Table 2 did not parse (%d cells)" % len(out)
        return out
    return cached(("dor", target), build)


def jobs(target):
    return cached(("jobs", target), lambda: {int(r["JOB_NO"]): r for r in rows_csv(p(target, "jobs"))})


def history(target):
    def build():
        out = {}
        for r in rows_csv(p(target, "history")):
            out.setdefault(int(r["JOB_NO"]), []).append(r)
        return out
    return cached(("hist", target), build)


def ledger(target):
    return cached(("ledger", target), lambda: rows_csv(p(target, "ledger")))


def previous_numbers(target):
    """The old system's job numbers, from the notes the conversion wrote into each moved job's history."""
    def build():
        out = {}
        for no, evs in history(target).items():
            for e in evs:
                m = re.match(r"Converted; previous job no\. (\d+)$", e["NOTE"])
                if e["STATUS"] == "MOVED" and m:
                    out[int(m.group(1))] = no
        return out
    return cached(("prev", target), build)


def day(s):
    return dt.date.fromisoformat(s[:10])


# ---------------------------------------------------------------------------------------------- the county counts
def shares(target, ratio="res", fallback=True, key="rel"):
    """ZIP code -> {county: weight}. The policy's rule: the residential ratio, and a ZIP code with no residential
    addresses by its total ratio. Wrong routes: another ratio, P.O. Box ZIP codes dropped, or the county-to-ZIP file
    read as if it gave each ZIP code's split."""
    rows = relationship(target, key)
    by = {}
    for r in rows:
        by.setdefault(r["zip"], []).append(r)
    out = {}
    for z, rs in by.items():
        use = ratio
        if ratio == "res" and sum(r["res"] for r in rs) == 0:
            if not fallback:
                continue
            use = "tot"
        out[z] = {r["fips"]: r[use] for r in rs if r[use] > 0}
    return out


def averages(d, classes=IRS_CLASS1_DOR):
    """The credit's average per childless return and per return with qualifying children in the IRS's first class
    (the Department's classes 1 to 3), and the share of that class's credit returns with children."""
    n0 = sum(d[(c, 0)][0] for c in classes)
    s0 = sum(d[(c, 0)][1] for c in classes) * 1000.0
    n1 = sum(d[(c, k)][0] for c in classes for k in (1, 2, 3))
    s1 = sum(d[(c, k)][1] for c in classes for k in (1, 2, 3)) * 1000.0
    return s0 / n0, s1 / n1, n1 / (n0 + n1)


def pooled_family_average(d):
    n = sum(v[0] for (c, k), v in d.items() if k > 0)
    s = sum(v[1] for (c, k), v in d.items() if k > 0) * 1000.0
    return s / n


def overall_family_share(d):
    n = sum(v[0] for v in d.values())
    return sum(v[0] for (c, k), v in d.items() if k > 0) / n


def cell_families(n, amt, a0, a1, route, share=None, pa1=None, oshare=None, scale=1000.0):
    """Returns with qualifying children in one ZIP code's first-class cell under each route."""
    if route == "R0":
        return n
    if route == "R1":
        return n * share
    if route == "R1all":
        return n * oshare
    if route == "R1b":
        return (amt * scale - a0 * n) / (pa1 - a0)
    if route == "R1c":
        return amt * scale / a1
    return (amt * scale - a0 * n) / (a1 - a0)


def families(target, route="gold", key="irs23", ratio="res", fallback=True, rel_key="rel", classes=IRS_CLASS1_DOR,
             scale=1000.0):
    """Each member county's returns claiming the credit with at least one qualifying child.
    gold   the first-class cell split into childless and family returns by the class's two average credits
    R0     every credit return
    R1     the Department's statewide share of family returns in the first class (the default path)
    R1all  the statewide share over every class, applied to every credit return
    R1b    the split with the family average pooled over every class
    R1c    the first class's credit amount over the family average alone
    Above the first class every credit return has a qualifying child (the Department's table shows none without)."""
    cells = irs(target, key)
    d = dor(target)
    a0, a1, share = averages(d, classes)
    pa1, osh = pooled_family_average(d), overall_family_share(d)
    sh = shares(target, ratio, fallback, rel_key)
    fips = member_fips(target)
    out = {f: 0.0 for f in fips}
    for (z, s), (n1, agi, ec, ea) in cells.items():
        if z in ("00000", "99999") or z not in sh:
            continue
        if route == "R1all":
            v = ec * osh
        elif s > 1 or route == "R0":
            v = ec
        else:
            v = cell_families(ec, ea, a0, a1, route, share, pa1, osh, scale)
        for f, w in sh[z].items():
            if f in out:
                out[f] += w * v
    return out


def allocate(counts, fund=FUND, cap=CAP):
    """Shares of the fourteen counties' total; a county above the cap held at it and the excess shared in proportion
    among the rest, repeated until none is above it. Unrounded."""
    fs = list(counts)
    v = np.array([counts[f] for f in fs], float)
    out = np.zeros(len(v))
    fixed = np.zeros(len(v), bool)
    while True:
        rest = fund - out[fixed].sum()
        free = ~fixed
        out[free] = rest * v[free] / v[free].sum()
        over = free & (out > cap * fund + 1e-9)
        if not over.any():
            break
        out[over] = cap * fund
        fixed |= over
    return {f: float(x) for f, x in zip(fs, out)}


def allocations(target, route="gold", **kw):
    return allocate(families(target, route, **kw))


def capped(al, fund=FUND, cap=CAP):
    return sorted(f for f, x in al.items() if abs(x - cap * fund) < 1e-6)


def county_table(target, route="gold", **kw):
    """Per county: returns with qualifying children, share of the fourteen counties' total (per cent), allocation."""
    fam = families(target, route, **kw)
    tot = sum(fam.values())
    al = allocate(fam)
    return {f: dict(families=fam[f], share=100.0 * fam[f] / tot, allocation=al[f]) for f in fam}


# ---------------------------------------------------------------------------------------------- ask A: dwellings
def first_completed(target):
    out = {}
    for no, evs in history(target).items():
        c = sorted(e["CHANGED_AT"] for e in evs if e["STATUS"] == "COMPLETED")
        if c:
            out[no] = day(c[0])
    return out


def completions(target):
    out = {}
    for no, evs in history(target).items():
        out[no] = sorted(day(e["CHANGED_AT"]) for e in evs if e["STATUS"] == "COMPLETED")
    return out


def tract_county(r):
    return r["TRACT"][:5]


def field_county(target, r):
    nm = {n: f for f, n in names(target).items()}
    return nm[r["COUNTY"]]


def program(r):
    return any(x in r["FUNDING"].split(";") for x in PROGRAM)


def dwellings(target, route="gold"):
    """Dwellings the program weatherized in each member county in program year 2025.
    gold            units of DOE or state funded jobs first completed in the year (the status history), by the
                    dwelling's tract
    as_exported     the export's completion date (the latest) in place of the first
    as_exported_cf  the export's completion date and the county the agency entered
    county_field    the first completion, the county the agency entered
    jobs            jobs counted in place of units
    with_util       utility-only jobs counted too
    any_completion  any completion in the year, a callback's included
    current         jobs whose status is COMPLETED and whose export date falls in the year"""
    js = jobs(target)
    fc = first_completed(target)
    comps = completions(target)
    out = {f: 0 for f in member_fips(target)}
    for no, r in js.items():
        if route != "with_util" and not program(r):
            continue
        if route in ("as_exported", "as_exported_cf"):
            ok = bool(r["COMPLETED_DATE"]) and PY0 <= day(r["COMPLETED_DATE"]) <= PY1
        elif route == "current":
            ok = r["STATUS"] == "COMPLETED" and PY0 <= day(r["COMPLETED_DATE"]) <= PY1
        elif route == "any_completion":
            ok = any(PY0 <= c <= PY1 for c in comps.get(no, []))
        else:
            ok = no in fc and PY0 <= fc[no] <= PY1
        if not ok:
            continue
        f = field_county(target, r) if route in ("as_exported_cf", "county_field") else tract_county(r)
        out[f] += 1 if route == "jobs" else int(r["UNITS"])
    return out


# ---------------------------------------------------------------------------------------------- ask B: spending
def job_of_line(target, ln, how="gold"):
    """The job a ledger line is charged to. The invoices gave the old system's number until the move and the WX
    number after it; the conversion notes carry the old numbers. raw: the line's number joined as written."""
    if not ln["JOB_NO"]:
        return None
    no = int(ln["JOB_NO"])
    if how == "raw":
        return no if no in jobs(target) else None
    prev = previous_numbers(target)
    js = jobs(target)
    if day(ln["POSTED_DATE"]) < MOVE:
        return prev.get(no)
    # after the move: the WX number, except an invoice prepared before it; the agency and the service date decide
    cands = [c for c in (no, prev.get(no)) if c in js]
    good = [c for c in cands if js[c]["AGENCY"] == ln["AGENCY"] and js[c]["AUDIT_DATE"] <= ln["SERVICE_DATE"]]
    return good[0] if len(good) == 1 else (cands[0] if cands else None)


def spending(target, route="gold"):
    """The program's spending on dwellings in each member county in program year 2025.
    gold            ledger lines posted in the year and charged to a job, at the amount posted, placed by the job's
                    dwelling tract, each line's job read through the old number it was invoiced under
    raw_join        the line's job number joined to the export as written
    raw_join_cf     as written, placed by the county the agency entered
    county_field    the right job, the county the agency entered
    service_year    lines whose service date falls in the year in place of the posting date
    completed_only  only lines on jobs completed (weatherized) by the pull"""
    js = jobs(target)
    out = {f: 0.0 for f in member_fips(target)}
    for ln in ledger(target):
        if not ln["JOB_NO"]:
            continue
        d = day(ln["SERVICE_DATE"] if route == "service_year" else ln["POSTED_DATE"])
        if not (PY0 <= d <= PY1):
            continue
        no = job_of_line(target, ln, "raw" if route.startswith("raw_join") else "gold")
        if no is None:
            continue
        r = js[no]
        if route == "completed_only" and not r["COMPLETED_DATE"]:
            continue
        f = field_county(target, r) if route in ("raw_join_cf", "county_field") else tract_county(r)
        out[f] += float(ln["AMOUNT"])
    return out


A_ROUTES = ["as_exported", "as_exported_cf", "county_field", "jobs", "with_util", "any_completion", "current"]
B_ROUTES = ["raw_join", "raw_join_cf", "county_field", "service_year", "completed_only"]
MAIN_ROUTES = ["R0", "R1", "R1all", "R1b", "R1c"]
