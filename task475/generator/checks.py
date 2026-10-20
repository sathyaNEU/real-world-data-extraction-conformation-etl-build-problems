"""The build's assertions, on the cut pack and the clean world: the input gates, the realism sweep, the semantic fork
table (SF1 to SF14), the ladder, the oracles, the three rules (R1 the raw material ships, R2 each Layer 1 item
mishandled leaves the call, R3a the easy route on the clean truth, R3b the decisive step alone, R3c no plant),
separation, the two asks with their devices and hazards, the edges, figures.json and the goldens."""
import csv
import datetime as dt
import json
import os
import re
import shutil
import tempfile
from collections import Counter

import numpy as np
import openpyxl

import bundle as B
import pack as P
import world as W

CALL_FILES = ("irs23", "irs22", "rel", "rel_rev", "dor", "policy", "members", "agencies", "dictionary", "manifest")
ASK_FILES = ("members", "agencies", "jobs", "history", "ledger", "closeout", "dictionary", "manifest", "policy")


def r100(v):
    return int(round(v / 100.0)) * 100


def run(target, gdir=None, x=None, meta_path=None, fig_path=None, quiet=False):
    x = x or W.build()
    fails = []

    def ok(cid, cond, msg):
        if not quiet or not cond:
            print("%-5s %s  %s" % (cid, "ok  " if cond else "FAIL", msg))
        if not cond:
            fails.append(cid)

    nm = B.names(target)
    fips = B.member_fips(target)
    files = sorted(os.listdir(target))
    texts = {}
    for f in files:
        path = os.path.join(target, f)
        if f.endswith(".pdf"):
            texts[f] = B.pdf_text(path)
        elif f.endswith((".csv", ".txt")):
            texts[f] = open(path, encoding="utf-8").read()
    gold = B.county_table(target, "gold")
    al_g = {f: v["allocation"] for f, v in gold.items()}

    # ------------------------------------------------------------------ input gates
    ok("G1", len(files) >= 10 and set(files) == set(P.FILES.values()), "%d files, the declared set" % len(files))
    fmts = Counter(f.rsplit(".", 1)[1] for f in files)
    ok("G2", len(fmts) >= 3, "formats %s" % dict(fmts))
    nrows = {}
    for f in files:
        if f.endswith(".csv"):
            with open(os.path.join(target, f), encoding="utf-8", newline="") as fh:
                nrows[f] = sum(1 for _ in csv.reader(fh)) - 1
        elif f.endswith(".xlsx"):
            wb = openpyxl.load_workbook(os.path.join(target, f), read_only=True)
            nrows[f] = sum(1 for _ in wb[wb.sheetnames[0]].iter_rows()) - 1
            wb.close()
    big = max(nrows, key=nrows.get)
    ok("G3", nrows[big] >= 25000, "largest file %s, %d rows" % (big, nrows[big]))
    if meta_path and os.path.exists(meta_path):
        meta = json.load(open(meta_path))
        ok("G4", len(meta["distractor_files"]) >= 2 and all(d in files for d in meta["distractor_files"])
           and all(meta["rows"].get(k) == v for k, v in nrows.items() if k in meta["rows"]),
           "distractors %s; metadata row counts agree" % meta["distractor_files"])
    man = B.rows_csv(os.path.join(target, P.FILES["manifest"]))
    ok("G5", {r["FILE"] for r in man} == set(files) - {P.FILES["manifest"]},
       "the manifest lists every other shipped file")
    ok("G6", all(str(nrows[r["FILE"]]) == r["ROWS"] for r in man if r["FILE"] in nrows),
       "manifest row counts equal the files")
    dic = texts[P.FILES["dictionary"]]
    ok("G7", all(f in dic for f in files if f != P.FILES["dictionary"]), "the dictionary names every other file")

    # ------------------------------------------------------------------ dates (no record after its extract)
    js = B.jobs(target)
    hist = B.history(target)
    gl = B.ledger(target)
    dmax = max(max(r[c] for c in ("OPENED", "AUDIT_DATE", "START_DATE", "COMPLETED_DATE") if r[c]) for r in js.values())
    hmax = max(e["CHANGED_AT"] for evs in hist.values() for e in evs)
    lmax = max(r["POSTED_DATE"] for r in gl)
    early = sum(1 for r in gl if r["SERVICE_DATE"] > r["POSTED_DATE"])
    ok("D1", dmax <= "2026-09-30" and hmax <= "2026-09-30 23:59" and lmax <= "2026-09-30" and early == 0,
       "export to %s, history to %s, ledger posted to %s, no service after posting" % (dmax, hmax, lmax))
    ok("D2", min(r["POSTED_DATE"] for r in gl) >= "2024-07-01", "ledger from %s" % min(r["POSTED_DATE"] for r in gl))

    # ------------------------------------------------------------------ the policy's filed sentences
    pol = re.sub(r"\s+", " ", texts[P.FILES["policy"]])
    pins = ["at least one qualifying child", "latest tax year", "residential addresses", "total addresses",
            "counted for no county", "12 per cent", "repeated until no county is above 12 per cent",
            "nearest hundred dollars", "first completed", "close-out counts spending"]
    ok("P1", all(p_ in pol for p_ in pins), "the policy carries %d pins" % len(pins))
    co = re.sub(r"\s+", " ", texts[P.FILES["closeout"]])
    ok("P2", "posts the line, at the amount posted" in co and "charged to jobs" in co
       and "pays for" in co, "the close-out carries the spending and program-job conventions")
    vocab = ["childless", "decompos", "average credit", "per return", "renumber", "legacy", "mix of"]
    hits = {v: [f for f, t in texts.items() if v in t.lower()] for v in vocab}
    ok("P3", not any(hits.values()), "no shipped text uses the decisive or device vocabulary: %s" % {
        k: v for k, v in hits.items() if v})
    co_bare = co
    for a_ in B.rows_csv(os.path.join(target, P.FILES["agencies"])):
        co_bare = co_bare.replace(a_["AGENCY_NAME"], "")
    ok("P4", not any(nm[f] in co_bare for f in fips), "the close-out names no county outside the agencies' names")

    # ------------------------------------------------------------------ the fork table
    def table(**kw):
        return B.county_table(target, **kw)

    def capped_names(t):
        return [nm[f] for f in B.capped({k: v["allocation"] for k, v in t.items()})]

    def misses(t):
        return [nm[f] for f in fips if r100(t[f]["allocation"]) != r100(gold[f]["allocation"])]

    r0 = table(route="R0")
    ok("SF1", capped_names(r0) == ["Tallis"] and not capped_names(gold),
       "every credit return: Tallis %.2f per cent, capped; the policy's returns with children: none capped" % (
           r0["57005"]["share"]))
    t22 = table(key="irs22")
    ok("SF2", len(misses(t22)) == 14 and "TAX_YEAR" not in texts[P.FILES["irs22"]][:200]
       and texts[P.FILES["irs22"]].startswith("statefips"), "the 2022 file moves all 14 allocations; its "
                                                            "lower-case header and the dictionary date it")
    tt = table(ratio="tot")
    trev = table(rel_key="rel_rev")
    ok("SF3", not capped_names(tt) and len(misses(tt)) >= 12 and capped_names(trev) != [],
       "the total ratio moves %d allocations, same call; the county-to-ZIP file read as a split caps %s" % (
           len(misses(tt)), capped_names(trev)))
    rel = B.relationship(target)
    pobox = sorted({r["zip"] for r in rel} - {z for z in {r["zip"] for r in rel}
                                              if sum(q["res"] for q in rel if q["zip"] == z) > 0})
    tnp = table(fallback=False)
    cells = B.irs(target)
    ok("SF4", len(pobox) == 4 and all((z, 1) in cells for z in pobox) and not capped_names(tnp)
       and len(misses(tnp)) >= 10, "%d P.O. Box ZIP codes with returns and no residential addresses; dropped, "
                                   "%d allocations move, same call" % (len(pobox), len(misses(tnp))))
    relz = {r["zip"] for r in rel}
    ok("SF5", "00000" not in relz and "99999" not in relz and ("99999", 1) in cells,
       "the state and small-ZIP lines are in the IRS file and in no relationship row")
    counties_rel = {r["fips"] for r in rel}
    ok("SF6", len(fips) == 14 and len(counties_rel) > 14 and set(fips) <= counties_rel
       and abs(sum(v["share"] for v in gold.values()) - 100.0) < 1e-9,
       "14 members of %d counties in the relationship file; shares of the fourteen" % len(counties_rel))
    dol = B.averages(B.dor(target))
    tdol = {f: 0.0 for f in fips}
    for (z, s), (n1, agi, ec, ea) in B.irs(target).items():
        if s == 1 and z in B.shares(target):
            for f, w in B.shares(target)[z].items():
                if f in tdol:
                    tdol[f] += w * (ea - dol[0] * ec) / (dol[1] - dol[0])
    ok("SF7", all(v < 0 for v in tdol.values()), "amounts read as dollars: every county's first-class count of "
                                                  "returns with children negative")
    d = B.dor(target)
    ok("SF8", all(d[(c, 0)] == (0, 0) for c in (4, 5, 6)) and all(d[(c, 0)][0] > 0 for c in (1, 2, 3)),
       "childless credit returns only below $25,000 (the IRS's first class)")
    ok("SF9", all(B.cell_families(1, 1, 0, 1, "gold") is not None for _ in [0]) and
       all(d[(c, 0)][0] == 0 for c in (4, 5, 6)), "classes above the first are all returns with children")
    a0, a1, share = B.averages(d)
    sh = B.shares(target)
    agg = {f: [0.0, 0.0, 0.0] for f in fips}      # class-1 claims, class-1 amount, claims above class 1
    implied = []
    for (z, s), (n1, agi, ec, ea) in cells.items():
        if z in ("00000", "99999") or z not in sh:
            continue
        if s == 1 and ec:
            implied.append(((ea * 1000.0 - a0 * ec) / (a1 - a0)) / ec)
        for f, w in sh[z].items():
            if f in agg:
                if s == 1:
                    agg[f][0] += w * ec
                    agg[f][1] += w * ea
                else:
                    agg[f][2] += w * ec
    cty_first = {f: (v[1] * 1000.0 - a0 * v[0]) / (a1 - a0) + v[2] for f, v in agg.items()}
    ok("SF10", max(abs(cty_first[f] - gold[f]["families"]) for f in fips) < 1e-6 and min(implied) > 0
       and max(implied) < 1, "the split by ZIP code then county equals the split by county; every cell's share "
                             "of returns with children in (0, 1): %.3f to %.3f" % (min(implied), max(implied)))
    t1b = table(route="R1b")
    st = {s: cells[("00000", s)] for s in range(1, 7)}
    fam_state_gold = (st[1][3] * 1000.0 - a0 * st[1][2]) / (a1 - a0)
    pa1 = B.pooled_family_average(d)
    fam_state_1b = (st[1][3] * 1000.0 - a0 * st[1][2]) / (pa1 - a0)
    dor_fam1 = sum(d[(c, k)][0] for c in (1, 2, 3) for k in (1, 2, 3))
    ok("SF11", not capped_names(t1b) and len(misses(t1b)) == 14 and abs(fam_state_gold - dor_fam1) < 25
       and abs(fam_state_1b - dor_fam1) > 500,
       "pooled family average: no cap, all 14 allocations miss; on the state line the class averages give %.0f "
       "returns with children in the first class against the Department's %d, the pooled average %.0f" % (
           fam_state_gold, dor_fam1, fam_state_1b))
    ok("SF12", not capped_names(gold) and max(v["share"] for v in gold.values()) < 11.9,
       "no county reaches the cap, so one pass and the repeated rule agree")
    fr = {f: round(gold[f]["families"], 1) for f in fips}
    alr = B.allocate(fr)
    tot_r = sum(fr.values())
    al_sh = {f: round(100.0 * fr[f] / tot_r, 2) / 100.0 * B.FUND for f in fips}
    ok("SF13", all(r100(alr[f]) == r100(al_g[f]) for f in fips)
       and sum(r100(al_sh[f]) != r100(al_g[f]) for f in fips) >= 10,
       "counts rounded to one decimal first: the same allocations; shares rounded first: %d move" % (
           sum(r100(al_sh[f]) != r100(al_g[f]) for f in fips)))
    ok("SF14", abs(sum(al_g.values()) - B.FUND) < 1e-6, "the fund is shared whole: $%s" % "{:,.2f}".format(
        sum(al_g.values())))

    # ------------------------------------------------------------------ the ladder
    r1 = table(route="R1")
    r1c = table(route="R1c")
    ok("L1", capped_names(r0) == ["Tallis"], "R0 every credit return: Tallis %.2f per cent, capped" % r0["57005"][
        "share"])
    ok("L2", capped_names(r1) == ["Tallis"] and r1["57005"]["share"] > 13.5,
       "R1 the default path, the statewide share of returns with children: Tallis %.2f per cent, capped" % r1[
           "57005"]["share"])
    ok("L3", not capped_names(gold) and gold["57005"]["share"] < 11.75
       and max(gold, key=lambda f: gold[f]["share"]) == "57005",
       "the golden: Tallis the largest at %.2f per cent, under the cap" % gold["57005"]["share"])
    wrong = {"R0": r0, "R1": r1, "R1b": t1b, "R1c": r1c}
    ok("L4", all(len(misses(t)) == 14 for t in wrong.values()), "every wrong route misses all 14 allocations: %s" % {
        k: len(misses(t)) for k, t in wrong.items()})
    second = {k: sorted(v["share"] for v in t.values())[-2] for k, t in list(wrong.items()) + [("gold", gold)]}
    ok("L5", all(v < 10.5 for v in second.values()), "the second-largest share under every route: %s" % {
        k: round(v, 2) for k, v in second.items()})
    dm = [abs(r1[f]["allocation"] - gold[f]["allocation"]) for f in fips]
    ok("L6", min(dm) >= 1500, "the default path and the golden differ by $%s to $%s per county" % (
        "{:,.0f}".format(min(dm)), "{:,.0f}".format(max(dm))))
    gap1b = [abs(t1b[f]["allocation"] - gold[f]["allocation"]) for f in fips]
    ok("L7", min(gap1b) >= 150, "the pooled average differs by $%s to $%s per county" % (
        "{:,.0f}".format(min(gap1b)), "{:,.0f}".format(max(gap1b))))
    ok("L8", capped_names(r1c) == ["Tallis"], "the class amount over the family average alone: Tallis %.2f per "
                                               "cent, capped" % r1c["57005"]["share"])

    # ------------------------------------------------------------------ oracles
    zl = {s: [0, 0, 0, 0] for s in range(1, 7)}
    for (z, s), v in cells.items():
        if z != "00000":
            for i in range(4):
                zl[s][i] += v[i]
    ok("O1", all(zl[s][0] == st[s][0] and zl[s][2] == st[s][2] for s in range(1, 7))
       and all(abs(zl[s][3] - st[s][3]) <= 0.5 * sum(1 for (z, q) in cells if q == s and z != "00000")
               for s in range(1, 7)), "the state line ties the ZIP lines and the small-ZIP line in every class")
    dor_claims = sum(v[0] for v in d.values())
    ok("O2", dor_claims == sum(st[s][2] for s in range(1, 7)) and sum(d[(c, k)][0] for c in (1, 2, 3)
                                                                       for k in range(4)) == st[1][2],
       "the Department's credit returns equal the IRS state line's (%d), the first class's too" % dor_claims)
    pub_amt = {f: agg[f][1] * 1000.0 for f in fips}
    r1_amt = {f: agg[f][0] * (share * a1 + (1 - share) * a0) for f in fips}
    errs = {nm[f]: 100.0 * (r1_amt[f] / pub_amt[f] - 1) for f in fips}
    ok("O3", max(abs(v) for v in errs.values()) > 8 and errs["Tallis"] > 15,
       "the statewide share implies first-class credit amounts the counties' own amounts contradict: Tallis "
       "%+.1f per cent, range %+.1f to %+.1f" % (errs["Tallis"], min(errs.values()), max(errs.values())))

    # ------------------------------------------------------------------ the three rules
    cols_all = set()
    for f in files:
        if f.endswith(".csv"):
            cols_all |= set(texts[f].splitlines()[0].upper().split(","))
    fam_re = [re.compile(r"(?<![\d.,])%s(?![\d])" % re.escape("%.1f" % gold[f]["families"])) for f in fips]
    fam_hits = [f_ for f_, t in texts.items() for rx in fam_re if rx.search(t)]
    ok("R1", {"N59660", "A59660"} <= cols_all and not fam_hits
       and len(d) == 24, "the raw material ships (ZIP x class credit returns and amounts, the Department's table by "
                         "children); no shipped file holds a county count of returns with children")
    # the clean truth: the world's cells and the Department's table unrounded
    ty = x["ty"]
    clean = {k: (v[0], v[1] / 1000.0, v[2], v[3] / 1000.0) for k, v in ty["raw"].items()}
    clean_dor = {k: (v[0], v[1] / 1000.0) for k, v in ty["raw_dor"].items()}
    for c in range(1, 7):
        for k in range(4):
            clean_dor.setdefault((c, k), (0, 0.0))
    wsh = x["sh"]
    def w_alloc(route):
        fam = W.families(clean, clean_dor, wsh, route)
        al = W.allocate(fam)
        return fam, al
    fam_c1, al_c1 = w_alloc("R1")
    fam_cg, al_cg = w_alloc("gold")
    truth = {f: 0 for f in fips}
    for z, f, kids, mfj, a, cr in ty["claim"]:
        if kids > 0 and f in truth:
            truth[f] += 1
    tsh = 100.0 * truth["57005"] / sum(truth.values())
    ok("R3a", W.capped(al_c1) == ["Tallis"] and W.capped(al_cg) == ["none"] and tsh < 12.0,
       "the easy route on the clean truth caps Tallis (%.2f per cent); the split leaves it at %.2f; the true count "
       "of families puts it at %.2f" % (100 * fam_c1["57005"] / sum(fam_c1.values()),
                                        100 * fam_cg["57005"] / sum(fam_cg.values()), tsh))
    ok("R3b", capped_names(r0) == ["Tallis"] and capped_names(r1) == ["Tallis"] and not capped_names(gold),
       "R0 to R1 (the statewide share added) still caps Tallis; R1 with only the split changed is the golden")
    co_text = texts[P.FILES["closeout"]]
    ok("R3c", not any(("%s" % "{:,}".format(r100(al_g[f]))) in t for t in texts.values() for f in fips)
       and "Allocation" not in co_text and len([f for f in files if "memo" in f or "draft" in f]) == 0,
       "no shipped file states an allocation, ranks the counties or argues a basis; nothing planted")
    # R2: each Layer 1 item mishandled leaves the call
    ok("R2", True, "Layer 1 items, each mishandled:")
    up = [c.upper() for c in texts[P.FILES["irs22"]].splitlines()[0].split(",")]
    ok("R2a", up == P.IRS23_COLS, "the 2022 release's lower-case layout is the 2023 layout by position")
    ok("R2b", all(v < 0 for v in tdol.values()), "amounts in thousands read as dollars: the split fails loudly "
                                                  "(every county negative); the call cannot be filed from it")
    ok("R2c", "00000" not in relz and "99999" not in relz, "the state and small-ZIP lines fall out of any ZIP to "
                                                           "county join; the call unchanged")
    ok("R2d", not capped_names(tnp), "the P.O. Box ZIP codes dropped: no county capped (%d allocations move, each "
                                     "ruled out by policy 4.1's fallback)" % len(misses(tnp)))
    ok("R2e", not capped_names(tt), "the total ratio in place of the residential: no county capped (%d move, ruled "
                                    "out by policy 4.1)" % len(misses(tt)))
    rowsum = Counter()
    for r in rel:
        rowsum[r["zip"]] += r["res"]
    ok("R2f", all(abs(v - 1.0) < 1e-5 or v == 0 for v in rowsum.values()), "residential ratios sum to one per ZIP "
                                                                            "code (or zero for P.O. Box ZIP codes)")

    # ------------------------------------------------------------------ separation
    def subset(keys, fn):
        tmp = tempfile.mkdtemp()
        try:
            for k in keys:
                shutil.copy(os.path.join(target, P.FILES[k]), tmp)
            return fn(tmp)
        finally:
            shutil.rmtree(tmp)
    sg = subset(CALL_FILES, lambda t: B.county_table(t, "gold"))
    ok("S1", all(abs(sg[f]["allocation"] - gold[f]["allocation"]) < 1e-9 for f in fips),
       "the call unchanged with the job export, history, ledger and close-out removed")
    ga, gb = B.dwellings(target), B.spending(target)
    sa = subset(ASK_FILES, lambda t: (B.dwellings(t), B.spending(t)))
    ok("S2", sa[0] == ga and all(abs(sa[1][f] - gb[f]) < 1e-6 for f in fips),
       "the asks unchanged with the IRS, relationship and Department files removed")

    # ------------------------------------------------------------------ ask A: dwellings
    at = W.a_truth(x["jobs"])
    ok("A1", at == ga, "ask A from the pack equals the clean truth: %s" % {nm[f]: ga[f] for f in fips})
    for cid, route, want in (("A2", "as_exported", 14), ("A3", "as_exported_cf", 14), ("A4", "current", 14),
                             ("A5", "jobs", 14), ("A6", "with_util", 13), ("A7", "any_completion", 8),
                             ("A15", "label", 14)):
        v = B.dwellings(target, route)
        m = [nm[f] for f in fips if v[f] != ga[f]]
        ok(cid, len(m) >= want, "%s misses %d: %s" % (route, len(m), {nm[f]: v[f] - ga[f] for f in fips}))
    v = B.dwellings(target, "county_field")
    ok("A8", sorted(nm[f] for f in fips if v[f] != ga[f]) == ["Hesketh", "Tallis"],
       "the county the agency entered moves Tallis and Hesketh")
    fc = B.first_completed(target)
    comps = B.completions(target)
    cross = Counter()
    for no, r in js.items():
        if B.program(r) and no in fc and r["COMPLETED_DATE"]:
            fi = B.PY0 <= fc[no] <= B.PY1
            la = B.PY0 <= B.day(r["COMPLETED_DATE"]) <= B.PY1
            if fi != la:
                cross[B.tract_county(r)] += 1
    ok("A9", all(cross[f] >= 1 for f in fips), "jobs whose first and latest completions straddle the year in every "
                                               "county: %s" % {nm[f]: cross[f] for f in fips})
    latest_ok = all(r["COMPLETED_DATE"] == (comps[no][-1].isoformat() if comps.get(no) else "")
                    for no, r in js.items() if r["STATUS"] in ("COMPLETED", "REOPENED"))
    twice = sum(1 for no in js if len(comps.get(no, [])) >= 2)
    ok("A10", latest_ok and twice >= 100, "the export's completion date is the latest COMPLETED in the history; %d "
                                          "jobs completed twice or more (the grader's check)" % twice)
    thca_h = sum(1 for r in js.values() if r["AGENCY"] == "THCA" and r["TRACT"][:5] == "57011" and r["COUNTY"] ==
                 "Tallis")
    ok("A11", thca_h >= 100 and all(r["COUNTY"] == nm[r["TRACT"][:5]] for r in js.values() if r["AGENCY"] != "THCA"),
       "the agency serving Tallis and Hesketh enters Tallis on its %d Hesketh jobs; every other agency the "
       "dwelling's county" % thca_h)
    mf = sum(1 for r in js.values() if int(r["UNITS"]) > 1)
    ok("A12", mf >= 300, "%d multifamily jobs, units over jobs" % mf)
    paid = B.paid_jobs(target)
    true_fund = {j["no"]: j["funding"] for j in x["jobs"]}
    hidden = Counter(B.tract_county(r) for no, r in js.items() if r["FUNDING"] == "UTIL-LIW" and no in paid
                     and no in fc and B.PY0 <= fc[no] <= B.PY1)
    ok("A13", all(hidden[f] >= 1 for f in fips), "dwellings the program paid for whose old record shows only the "
                                                 "utility's code, first completed in 2025, in every county: %s" % {
                                                     nm[f]: hidden[f] for f in fips})
    util_label = [no for no, r in js.items() if r["FUNDING"] == "UTIL-LIW"]
    util25 = [no for no in util_label if no in fc and B.PY0 <= fc[no] <= B.PY1]
    ok("A14", all((no in paid) == ("DOE-WAP" in true_fund[no]) for no in util25)
       and all(no in paid for no, r in js.items() if B.program(r) and no in fc and B.PY0 <= fc[no] <= B.PY1),
       "among jobs first completed in 2025 the ledger charges program money to every job the program paid for and to "
       "none of the utility's own %d" % sum(1 for no in util25 if no not in paid))
    on_util = sum(1 for ln in gl if ln["JOB_NO"] and B.job_of_line(target, ln) in set(util_label))
    combo_new = sum(1 for r in js.values() if r["FUNDING"] == "DOE-WAP;UTIL-LIW")
    ok("A16", on_util >= 50 and combo_new >= 10, "the grader's check: %d ledger lines of program money on jobs whose "
                                                 "record shows only UTIL-LIW; %d WX jobs list both codes" % (
                                                     on_util, combo_new))

    # ------------------------------------------------------------------ ask B: spending
    bt = W.b_truth(x["ledger"])
    ok("B1", all(abs(bt[f] - gb[f]) < 1e-6 for f in fips), "ask B from the pack equals the clean truth")
    for cid, route, want in (("B2", "raw_join", 14), ("B3", "raw_join_cf", 14), ("B4", "service_year", 13),
                             ("B5", "completed_only", 14), ("B11", "posted_cutoff", 14), ("B12", "service_cutoff", 14)):
        v = B.spending(target, route)
        m = [nm[f] for f in fips if round(v[f]) != round(gb[f])]
        ok(cid, len(m) >= want, "%s misses %d: %s" % (route, len(m), {nm[f]: round(v[f] - gb[f]) for f in fips}))
    v = B.spending(target, "county_field")
    ok("B6", sorted(nm[f] for f in fips if round(v[f]) != round(gb[f])) == ["Hesketh", "Tallis"],
       "the county the agency entered moves Tallis and Hesketh")
    prev = B.previous_numbers(target)
    l25 = [r for r in gl if r["JOB_NO"] and B.PY0 <= B.day(r["POSTED_DATE"]) <= B.PY1]

    def consistent(how):
        bad = 0
        for r in l25:
            no = B.job_of_line(target, r, how)
            if no is None or js[no]["AGENCY"] != r["AGENCY"] or r["SERVICE_DATE"] < js[no]["OPENED"]:
                bad += 1
        return bad
    early = sum(1 for r in l25 if B.invoice_month(r) < dt.date(2025, 7, 1))
    late_old = sum(1 for r in l25 if B.invoice_month(r) < dt.date(2025, 7, 1) and B.day(r["POSTED_DATE"]) >= B.MOVE)
    bads = {h: consistent(h) for h in ("gold", "posted_cutoff", "service_cutoff", "raw")}
    raw_unmatched = sum(1 for r in l25 if int(r["JOB_NO"]) not in js)
    ok("B7", bads["gold"] == 0 and all(v > 0 for k, v in bads.items() if k != "gold") and raw_unmatched == 0
       and late_old >= 100 and all(int(r["JOB_NO"]) in prev for r in l25 if B.invoice_month(r) < dt.date(2025, 7, 1)),
       "%d of %d 2025 lines are on invoices for service months to June, each with an old number (%d of them posted "
       "after the move); read by the invoice month every line's agency and dates agree; lines that disagree under the "
       "posting-date cutoff %d, the service-date cutoff %d, as written %d; none unmatched as written" % (
           early, len(l25), late_old, bads["posted_cutoff"], bads["service_cutoff"], bads["raw"]))
    moved = sum(1 for evs in hist.values() for e in evs if e["STATUS"] == "MOVED")
    noted = sum(1 for evs in hist.values() for e in evs if e["NOTE"] and e["STATUS"] != "MOVED")
    ok("B8", moved == len(prev) and noted >= moved, "%d conversion notes among %d other notes" % (moved, noted))
    raw = B.spending(target, "raw_join")
    ratio = {nm[f]: raw[f] / gb[f] for f in fips}
    ok("B9", all(0.4 < v < 1.8 for v in ratio.values()), "as written the county totals stay plausible: %s" % {
        k: round(v, 2) for k, v in ratio.items()})
    y25 = sum(float(r["AMOUNT"]) for r in gl if B.PY0 <= B.day(r["POSTED_DATE"]) <= B.PY1)
    co_tot = re.search(r"All spending ([\d,]+\.\d\d) ([\d,]+\.\d\d) ([\d,]+\.\d\d)", co)
    ok("B10", co_tot is not None and abs(float(co_tot.group(3).replace(",", "")) - y25) < 0.005,
       "the close-out's 2025 total ties the ledger's lines posted in 2025 under every handling: %s" % (
           co_tot.group(3) if co_tot else None))

    # ------------------------------------------------------------------ edges and ties
    far = []
    for f in fips:
        for val, unit, need, lab in ((gold[f]["families"], 0.1, 0.15, "families"), (gold[f]["share"], 0.01, 0.15,
                                                                                    "share"),
                                     (gold[f]["allocation"], 100.0, 0.2, "allocation"), (gb[f], 1.0, 0.15, "spending")):
            if not W.edge(val, unit, need):
                far.append((nm[f], lab, val))
    ok("E1", not far, "every graded figure off its rounding edge %s" % far)
    dup = {k: len(set(v)) for k, v in (("families", [round(gold[f]["families"], 1) for f in fips]),
                                       ("shares", [round(gold[f]["share"], 2) for f in fips]),
                                       ("allocations", [r100(al_g[f]) for f in fips]),
                                       ("dwellings", [ga[f] for f in fips]),
                                       ("spending", [round(gb[f]) for f in fips]))}
    ok("E2", all(v == 14 for v in dup.values()), "no two counties share a graded figure %s" % dup)
    tots = (round(sum(gold[f]["families"] for f in fips), 1), sum(ga.values()), round(sum(gb.values())))
    ok("E3", tots[1] % 10 != 0 and tots[2] % 100 != 0, "headline totals not round: %s" % (tots,))
    amt = Counter(r["AMOUNT"] for r in gl)
    ok("E4", amt.most_common(1)[0][1] <= 4, "no ledger amount repeats beyond chance (%s)" % (amt.most_common(1),))

    # ------------------------------------------------------------------ figures and goldens
    if fig_path and os.path.exists(fig_path):
        fig = json.load(open(fig_path))
        comp = fig["components"]
        ok("F1", all(comp["allocations"][nm[f]] == r100(al_g[f]) and comp["shares_pct"][nm[f]] ==
                     round(gold[f]["share"], 2) and comp["families"][nm[f]] == round(gold[f]["families"], 1)
                     for f in fips) and fig["call"]["capped"] == "none", "figures.json main block")
        ok("F2", all(fig["a_dwellings"][nm[f]] == ga[f] and fig["b_spending"][nm[f]] == round(gb[f]) for f in fips),
           "figures.json ask blocks")
        if gdir and os.path.isdir(gdir):
            import golden as GO
            wb = openpyxl.load_workbook(os.path.join(gdir, GO.WB))
            ws = wb["Allocation"]
            got = {}
            for row in ws.iter_rows(min_row=1, max_row=40, values_only=True):
                if row[0] in comp["allocations"]:
                    got[row[0]] = (row[3], row[4], row[5])
            ok("F3", all(got[k] == (comp["families"][k], comp["shares_pct"][k], comp["allocations"][k])
                         for k in comp["allocations"]), "the workbook's allocation table")
            ws2 = wb["Program year 2025"]
            got2 = {row[0]: (row[1], row[2]) for row in ws2.iter_rows(min_row=2, max_row=15, values_only=True)}
            ok("F4", all(got2[k] == (fig["a_dwellings"][k], fig["b_spending"][k]) for k in comp["allocations"]),
               "the workbook's program year 2025 table")
            with open(os.path.join(gdir, GO.CSV), encoding="utf-8") as fh:
                cr = list(csv.DictReader(fh))
            ok("F5", len(cr) == 14 and all(float(r["EITC_RETURNS_WITH_QUALIFYING_CHILDREN"]) == comp["families"][
                r["COUNTY"]] for r in cr), "the county table's 14 rows")
            people = set()
            for sh_ in wb.worksheets:
                for row in sh_.iter_rows(values_only=True):
                    for c in row:
                        if isinstance(c, str):
                            people |= {p_ for p_ in ("Brianna Wallace", "Mitchell Gonzalez") if p_ in c}
            ok("F6", people == {"Brianna Wallace", "Mitchell Gonzalez"}, "the golden's people are the drawn personas")
    print("checks: %d failed %s" % (len(fails), fails))
    return not fails


if __name__ == "__main__":
    import sys
    here = os.path.dirname(os.path.abspath(__file__))
    task = os.path.dirname(here)
    good = run(sys.argv[1] if len(sys.argv) > 1 else os.path.join(task, "target"), os.path.join(task, "golden"),
               meta_path=os.path.join(task, "metadata.json"), fig_path=os.path.join(here, "figures.json"))
    raise SystemExit(0 if good else 1)
