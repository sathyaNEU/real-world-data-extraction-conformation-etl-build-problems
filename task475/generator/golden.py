"""The golden deliverables, built from the shipped pack only (bundle.py): the 2027 allocation workbook (each member
county's tax-year 2023 credit returns with qualifying children, its share and its allocation, the chart against the
cap, and the program year 2025 record), the conformed county table, and generator/figures.json."""
import csv
import json
import os

import bundle as B

WB = "corran_valley_2027_weatherization_allocation.xlsx"
CSV = "corran_valley_ty2023_county_credit_families.csv"


def r100(v):
    return int(round(v / 100.0)) * 100


def compute(target):
    t = B.county_table(target, "gold")
    allc = B.families(target, "R0")
    nm = B.names(target)
    ag = {f: a for f, _, a in B.members(target)}
    order = B.member_fips(target)
    fam_tot = sum(v["families"] for v in t.values())
    top = max(order, key=lambda f: t[f]["share"])
    a0, a1, _ = B.averages(B.dor(target))
    riv = {k: B.county_table(target, k) for k in ("R1", "R0")}
    return dict(t=t, all=allc, nm=nm, ag=ag, order=order, fam_tot=fam_tot, top=top, a0=a0, a1=a1, riv=riv,
                capped=B.capped({f: v["allocation"] for f, v in t.items()}),
                a=B.dwellings(target), b=B.spending(target))


def figures(g):
    nm, t = g["nm"], g["t"]
    alloc = {nm[f]: r100(t[f]["allocation"]) for f in g["order"]}
    top = g["top"]
    return {
        "call": {"capped": "none" if not g["capped"] else ", ".join(nm[f] for f in g["capped"]),
                 "largest_share_pct": round(t[top]["share"], 2), "largest_allocation": alloc[nm[top]]},
        "components": {
            "largest_county": nm[top],
            "allocations": alloc,
            "shares_pct": {nm[f]: round(t[f]["share"], 2) for f in g["order"]},
            "families": {nm[f]: round(t[f]["families"], 1) for f in g["order"]},
            "credit_returns_all": {nm[f]: round(g["all"][f], 1) for f in g["order"]},
            "total_families": round(g["fam_tot"], 1),
            "allocations_total": sum(alloc.values()),
            "fund": B.FUND,
        },
        "a_dwellings": {nm[f]: int(g["a"][f]) for f in g["order"]},
        "b_spending": {nm[f]: int(round(g["b"][f])) for f in g["order"]},
        "ask_totals": {"dwellings": int(sum(g["a"].values())), "spending": int(round(sum(g["b"].values())))},
        "averages": {"childless": round(g["a0"], 2), "with_children": round(g["a1"], 2)},
        "rivals": {"statewide_share_largest_pct": round(g["riv"]["R1"][top]["share"], 2),
                   "every_credit_return_largest_pct": round(g["riv"]["R0"][top]["share"], 2)},
        "inputs": {"tax_year": 2023, "cap_pct": 12, "cap_dollars": int(B.CAP * B.FUND), "program_year": 2025,
                   "relationship_quarter": "2023 Q4",
                   "lines_for_no_county": ["00000", "99999"],
                   "policy": ["3.1", "4.1", "4.2", "5.1", "5.2", "5.3", "6.1"]},
    }


def write_csv(g, path):
    cols = ["TAX_YEAR", "COUNTY_FIPS", "COUNTY", "EITC_RETURNS", "EITC_RETURNS_WITH_QUALIFYING_CHILDREN"]
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(cols)
        for f in g["order"]:
            w.writerow([2023, f, g["nm"][f], "%.1f" % g["all"][f], "%.1f" % g["t"][f]["families"]])
    return len(g["order"])


def write_workbook(g, f, path):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.marker import DataPoint
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    nm, t, order = g["nm"], g["t"], g["order"]
    alloc = f["components"]["allocations"]
    top = g["top"]
    bold = Font(name="Calibri", bold=True)
    head = PatternFill("solid", fgColor="DDE7DF")
    thin = Side(style="thin", color="1B1F24")
    wb = Workbook()
    ws = wb.active
    ws.title = "Allocation"
    ws["A1"] = ("No county reaches the 12 per cent cap: %s County's %.2f per cent of the families claiming the credit "
                "is the largest share, $%s of the $20,000,000 fund" % (nm[top], t[top]["share"],
                                                                       "{:,}".format(alloc[nm[top]])))
    ws["A1"].font = Font(name="Calibri", bold=True, size=13, color="23413A")
    ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A1:F1")
    ws.row_dimensions[1].height = 36
    block = [("To", "Council Board, Corran Valley Council of Governments"),
             ("From", "Brianna Wallace, Director of Community Energy Programs"),
             ("Prepared by", "Mitchell Gonzalez, program analyst"),
             ("Date", "19 October 2026, for the 2027 weatherization plan due 15 November 2026"),
             ("Subject", "2027 weatherization fund: county allocation under the allocation policy, Revision 4")]
    r = 3
    for a, b in block:
        ws.cell(row=r, column=1, value=a).font = bold
        ws.cell(row=r, column=2, value=b)
        r += 1
    r += 1
    hdr = ["County", "FIPS", "Agency", "TY2023 credit returns with qualifying children",
           "Share of the fourteen counties (%)", "2027 allocation ($)"]
    h0 = r
    for k, v in enumerate(hdr, start=1):
        c = ws.cell(row=r, column=k, value=v)
        c.font = bold
        c.fill = head
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = 46
    r += 1
    first = r
    for fp in order:
        vals = [nm[fp], fp, g["ag"][fp], round(t[fp]["families"], 1), round(t[fp]["share"], 2), alloc[nm[fp]]]
        for k, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=k, value=v)
            if k == 4:
                c.number_format = "#,##0.0"
            elif k == 5:
                c.number_format = "0.00"
            elif k == 6:
                c.number_format = "#,##0"
        r += 1
    last = r - 1
    tot = ["Fourteen counties", None, None, round(g["fam_tot"], 1), 100.0, f["components"]["fund"]]
    for k, v in enumerate(tot, start=1):
        c = ws.cell(row=r, column=k, value=v)
        c.font = bold
        c.border = Border(top=thin)
        if k == 4:
            c.number_format = "#,##0.0"
        elif k == 5:
            c.number_format = "0.00"
        elif k == 6:
            c.number_format = "#,##0"
    r += 2
    note = ("Policy 5.2 holds a county at 12 per cent of the fund ($2,400,000); no county's share reaches it, so the "
            "fund is shared in proportion to the counts with nothing to redistribute. Counts and shares are carried "
            "unrounded and each allocation is rounded to the nearest hundred dollars (5.3); the rounded allocations "
            "sum to $%s, and the total shown is the fund." % "{:,}".format(f["components"]["allocations_total"]))
    ws.cell(row=r, column=1, value=note).font = Font(name="Calibri", italic=True, size=9, color="555555")
    ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    ws.row_dimensions[r].height = 40
    # the cap series for the chart, out of the print area
    ws.cell(row=h0, column=8, value="12 per cent cap ($)").font = bold
    ws.cell(row=h0, column=8).alignment = Alignment(wrap_text=True, vertical="center")
    for i in range(first, last + 1):
        ws.cell(row=i, column=8, value=int(B.CAP * B.FUND)).number_format = "#,##0"
    for col, wd in zip("ABCDEFGH", (14, 8, 9, 20, 16, 15, 3, 14)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = ws.cell(row=h0 + 1, column=2)
    ch = BarChart()
    ch.type = "col"
    ch.title = "2027 allocations against the 12 per cent cap"
    ch.y_axis.title = "Allocation ($)"
    ch.y_axis.number_format = "#,##0"
    ch.y_axis.majorGridlines = None
    data = Reference(ws, min_col=6, min_row=h0, max_row=last)
    cats = Reference(ws, min_col=1, min_row=first, max_row=last)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ser = ch.series[0]
    ser.graphicalProperties.solidFill = "8FA79B"
    ser.graphicalProperties.line.noFill = True
    hi = DataPoint(idx=order.index(top))
    hi.graphicalProperties.solidFill = "2C6E5A"
    hi.graphicalProperties.line.noFill = True
    ser.dPt = [hi]
    ser.dLbls = DataLabelList()
    ser.dLbls.showVal = True
    ser.dLbls.showSerName = False
    ser.dLbls.showCatName = False
    ser.dLbls.showLegendKey = False
    ser.dLbls.numFmt = "#,##0"
    ln = LineChart()
    ln.add_data(Reference(ws, min_col=8, min_row=h0, max_row=last), titles_from_data=True)
    ls = ln.series[0]
    ls.graphicalProperties.line.solidFill = "B23A2E"
    ls.graphicalProperties.line.dashStyle = "dash"
    ls.graphicalProperties.line.width = 19050
    ls.smooth = False
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = 2600000
    ch.y_axis.majorUnit = 500000
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch += ln
    ch.legend.position = "b"
    ch.gapWidth = 55
    ch.height, ch.width = 9.5, 15.6
    ws.add_chart(ch, "A%d" % (r + 2))

    ws2 = wb.create_sheet("Program year 2025")
    ws2.append(["County", "Dwellings weatherized in program year 2025", "Spending on dwellings in program year 2025 ($)"])
    for fp in order:
        ws2.append([nm[fp], f["a_dwellings"][nm[fp]], f["b_spending"][nm[fp]]])
    ws2.append(["Fourteen counties", f["ask_totals"]["dwellings"], f["ask_totals"]["spending"]])
    for c in ws2[1]:
        c.font = bold
        c.fill = head
        c.alignment = Alignment(wrap_text=True, vertical="center")
    lr = len(order) + 2
    for c in ws2[lr]:
        c.font = bold
        c.border = Border(top=thin)
    for row in ws2.iter_rows(min_row=2, max_row=lr, min_col=2, max_col=3):
        for c in row:
            c.number_format = "#,##0"
    ws2.row_dimensions[1].height = 34
    for col, wd in zip("ABC", (18, 22, 24)):
        ws2.column_dimensions[col].width = wd
    ws2.freeze_panes = "A2"
    src = ("Dwellings: units of the jobs the Department of Energy grant or the State Weatherization Fund paid for (the "
           "job's funding codes, or program money the grants ledger charges to it, since the previous system kept one "
           "code per job), in the year each job was first completed (WX job status history), placed by the dwelling's "
           "census tract. Spending: grants ledger lines posted in 2025 and charged to jobs, at the amount posted, each "
           "placed by its job's tract; invoices for service months to June 2025 were prepared before the July move "
           "and carry the previous system's job numbers, read through the conversion notes. Whole dollars from the "
           "unrounded county sums.")
    ws2.cell(row=lr + 2, column=1, value=src).font = Font(name="Calibri", italic=True, size=9, color="555555")
    ws2.cell(row=lr + 2, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws2.merge_cells(start_row=lr + 2, start_column=1, end_row=lr + 2, end_column=3)
    ws2.row_dimensions[lr + 2].height = 92

    ws3 = wb.create_sheet("Notes")
    notes = [
        ("Basis", "Policy 3.1: federal returns for tax year 2023 (the latest year of the IRS ZIP code statistics, "
                  "released 18 June 2026) claiming the earned income credit with at least one qualifying child, "
                  "filed by residents of each member county."),
        ("Residence", "Policy 4.1: each ZIP code's returns shared among counties by the residential ratio of the "
                      "State Geographic Information Office's ZIP code to county file for 2023 Q4; the four P.O. Box "
                      "ZIP codes, with no residential addresses, by their total ratio. The state line (00000) and "
                      "the combined small ZIP codes (99999) count for no county (4.2)."),
        ("Children", "The IRS statistics count credit returns without saying how many have qualifying children. "
                     "Above $25,000 every credit return has one (the Department of Revenue's Table 2 shows no "
                     "childless return above its $20,000 to $25,000 class). In the first IRS class, $1 under "
                     "$25,000 (the Department's first three classes), a ZIP code's returns with children are "
                     "(credit amount less the childless average credit times its credit returns) divided by the "
                     "difference between the family and childless average credits of that class statewide: "
                     "$%s and $%s." % ("{:,.2f}".format(g.get("a1", 0)), "{:,.2f}".format(g.get("a0", 0)))),
        ("Check", "The same split applied to the state line's first class gives the Department's count of returns "
                  "with children in its first three classes."),
        ("Cap", "Policy 5.2: 12 per cent of the fund, $2,400,000; the largest share is %s County's %.2f per cent, so "
                "the cap does not bind." % (nm[top], t[top]["share"])),
        ("Rounding", "Counts and shares unrounded; allocations to the nearest hundred dollars (5.3). The fourteen "
                     "rounded allocations sum to $%s." % "{:,}".format(f["components"]["allocations_total"])),
        ("Record", "Policy 6.1: the dwellings the program weatherized in each county in program year 2025 and the "
                   "program's spending on dwellings there, as the close-out counts spending (posted in the year)."),
        ("As of", "IRS statistics as released 18 June 2026; WX job system export and status history of 30 September "
                  "2026; grants ledger to 30 September 2026."),
    ]
    ws3.append(["Item", "Basis"])
    for a, b in notes:
        ws3.append([a, b])
    for c in ws3[1]:
        c.font = bold
        c.fill = head
    for row in ws3.iter_rows(min_row=2, max_row=len(notes) + 1):
        row[0].font = bold
        row[0].alignment = Alignment(vertical="top")
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    ws3.column_dimensions["A"].width = 14
    ws3.column_dimensions["B"].width = 100
    for i, (a, b) in enumerate(notes, start=2):
        ws3.row_dimensions[i].height = 16 * (1 + len(b) // 95)
    for sh in (ws, ws2, ws3):
        sh.page_setup.orientation = "landscape"
        sh.page_setup.paperSize = sh.PAPERSIZE_LETTER
        sh.sheet_properties.pageSetUpPr.fitToPage = True
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
    ws.print_area = "A1:F%d" % (last + 26)
    wb.save(path)


def main(target, gdir, fig_path):
    os.makedirs(gdir, exist_ok=True)
    g = compute(target)
    f = figures(g)
    write_workbook(g, f, os.path.join(gdir, WB))
    import postprocess as PP
    PP.normalize_ooxml(os.path.join(gdir, WB), "Mitchell Gonzalez", "Brianna Wallace", "2026-10-15T14:12:00Z",
                       "2026-10-19T21:06:33Z", excel=True, company="Corran Valley Council of Governments")
    write_csv(g, os.path.join(gdir, CSV))
    with open(fig_path, "w", encoding="utf-8") as fh:
        json.dump(f, fh, indent=1, sort_keys=False)
        fh.write("\n")
    return g, f
