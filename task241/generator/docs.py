"""task241 documents: the data management lead's release thread (eml), the reference data note (docx), the float
programme's model summary (pdf), the data dictionary and the readme (md)."""
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess

INSTITUTE = P.INSTITUTE
PEOPLE = {"dmqc": ("Naiara Nogueira", "nnogueira"), "lead": ("Tiago Tavares", "ttavares"),
          "eng": ("Patrícia Araújo", "paraujo")}
DOMAIN = "ioatlantico.org"


def _addr(k):
    name, user = PEOPLE[k]
    return "%s <%s@%s>" % (name, user, DOMAIN)


# ------------------------------------------------------------------ the release thread (eml)
def thread(path, F):
    lead, dmqc, eng = _addr("lead"), _addr("dmqc"), _addr("eng")
    body = f"""From: {lead}
To: {dmqc}
Cc: {eng}
Subject: RE: IBFA 2027 delayed-mode salinity release
Date: Fri, 7 Jan 2028 08:52:14 +0000
Message-ID: <c3e81f0a6d2b4e97@mail.{DOMAIN}>
In-Reply-To: <9b04d27e51a64c3f@mail.{DOMAIN}>
MIME-Version: 1.0
Content-Type: text/plain; charset="utf-8"
Content-Transfer-Encoding: 8bit

Naiara,

Yes to both. The partner institute's float stays out of everything in the release: we decode it, but it is not one of the array's 24. And in the delivery record a profile counts once, at the time its complete file reached us, so a profile whose first file broke off counts from the file that carried every level.

The reference data are as frozen on 15 December: our CTD archive and the RHD extract for the box, both in the folder with the reference data note. For the release report I need the same three tables as last year: each float's offset and what we do with it, the reference casts reaching 2,000 dbar for each year from 2015 to 2025, and the delivery record by month. The files go to the assembly centre on Friday 21 January and I would like to sign them off on the 19th.

Tiago Tavares
Data Management Lead, Ocean Data Centre
{INSTITUTE}, Lisboa

> From: {dmqc}
> Sent: Thursday, 6 January 2028 17:31
> To: {lead}
> Cc: {eng}
> Subject: RE: IBFA 2027 delayed-mode salinity release
>
> Tiago,
>
> The last 2027 files are in and Patrícia has the register and the technical files up to date. Two questions before I start. Does the partner float we decode go into the release or the delivery record? And for the delivery record, when a profile came in twice, the first file cut short, which time counts?
>
> Naiara
>
>> From: {lead}
>> Sent: Monday, 13 December 2027 10:05
>> To: {dmqc}
>> Cc: {eng}
>> Subject: IBFA 2027 delayed-mode salinity release
>>
>> Naiara, Patrícia,
>>
>> The 2027 delayed-mode release for the Iberian Basin Float Array closes the year's salinity record, so here is what goes into it.
>>
>> For each of the array's 24 floats, the 2027 salinity offset is the salinity the float reported less the true salinity of the water it sampled: the mean difference over the profile's good values between 1,400 and 1,900 dbar, averaged over the float's 2027 ascending profiles. Good values are the ones the real-time checks left at flag 1 or 2. A float whose offset is more than 0.010 in magnitude has its 2027 salinity adjusted in the release; every other float's 2027 salinity goes out as measured.
>>
>> The release report also carries the delivery record the assembly centre asks for: for each month of the year, how many of the array's ascending profiles reached us in full within 24 hours of the float surfacing, the profiles counted by the month the float surfaced.
>>
>> The reference data freeze on 15 December. Patrícia, could you bring the register and the cycle technical files up to date for the year?
>>
>> Tiago
"""
    with open(path, "w", encoding="utf-8", newline="\r\n") as fh:
        fh.write(body)


# ------------------------------------------------------------------ the reference data note (docx)
def reference_note(path, F):
    import docx
    from docx.shared import Pt
    doc = docx.Document()
    stl = doc.styles["Normal"]
    stl.font.name = "Calibri"
    stl.font.size = Pt(10.5)

    def para(t, bold_lead=None, after=6):
        p = doc.add_paragraph()
        if bold_lead:
            r = p.add_run(bold_lead + " ")
            r.bold = True
        p.add_run(t)
        p.paragraph_format.space_after = Pt(after)
        return p

    h = doc.add_paragraph()
    r = h.add_run("Reference data for delayed-mode salinity, Iberian Basin")
    r.bold = True
    r.font.size = Pt(13)
    for k, v in [("Prepared by", "Naiara Nogueira, Ocean Data Centre"), ("Frozen", "15 December 2027"),
                 ("For", "IBFA delayed-mode release, 2027 profiles")]:
        para(v, k + ":", after=0)
    doc.add_paragraph()
    para("The delayed-mode check compares each float's salinity with ship CTD data. The reference data for the 2027 "
         "release are frozen as of 15 December 2027 and come from two sources, each as its own pair of files: a station "
         "list and the levels.")
    para("Our own hydrographic cruises, 2015 to 2025, one a year, worked as a section across the basin. Files %s and "
         "%s. Each cast is processed to 10 dbar bins from the CTD's downcast; salinity is practical salinity, calibrated "
         "on each cruise against bottle samples run on the institute's salinometer. Temperature is in the IPTS-68 scale "
         "for the 2015 and 2016 cruises, which were processed before the processing software was replaced in 2017, and in "
         "ITS-90 from 2017; the temp_scale column says which. Pressure is in dbar, positions in decimal degrees, times in "
         "UTC. A cast stops where the station's planned depth or the bottom allowed, so the deepest level differs from "
         "cast to cast." % (F["arch_casts"], F["arch_levels"]), "Institute CTD archive.")
    para("Every CTD cast the RHD, the international reference hydrography database, holds for the box 36 to 44 N, 8 to "
         "22 W, from 2015 to 2025, as the database serves it, extracted on 15 December 2027. Files %s and %s. Levels "
         "are at the resolution each institute submitted. Depth is in metres (for casts submitted in pressure the "
         "database computes depth by Saunders 1981), temperature in ITS-90 to 0.001, practical salinity to 0.001; "
         "positions are given to 0.01 degree and the time of the cast as decimal hours UTC. Values the database's own quality control rejected are not in the "
         "extract." % (F["intl_casts"], F["intl_levels"]), "RHD extract.")
    para("Both sources' salinity is calibrated to the standard of the delayed-mode check, and no further correction is "
         "applied to either before use.", "Quality.")
    para("The float profiles for the year are in %s and %s; the float register is %s." %
         (F["profiles"], F["levels"], F["register"]), "Float data.")
    doc.core_properties.author = PEOPLE["dmqc"][0]
    doc.core_properties.title = "Reference data for delayed-mode salinity, Iberian Basin"
    doc.save(path)
    postprocess.scrub(path, PEOPLE["dmqc"][0], "2027-12-15")
    postprocess.finish_office(path, "2027-12-15T09:12:00Z", "2027-12-15T16:47:00Z", INSTITUTE)


# ------------------------------------------------------------------ the float models summary (pdf)
def models(path, reg):
    k = reg[reg["model"] == "K"]["wmo"].tolist()
    a = reg[reg["model"] == "A"]["wmo"].tolist()
    down = reg[reg["profiles"] == "ascending and descending"]["wmo"].tolist()
    first_a = min(reg[reg["model"] == "A"]["deployed"])
    import datetime as dt
    fa = dt.date.fromisoformat(first_a)
    rev = fa + dt.timedelta(days=12)
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.4, leading=13, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=14, leading=18, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=14,
                        spaceBefore=8, spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10.5, textColor=colors.HexColor("#444444"))
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=8.2, leading=10)
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                            bottomMargin=18 * mm, title="IBFA float models", author=PEOPLE["eng"][0],
                            creator=INSTITUTE, invariant=1)
    st = [Paragraph("Iberian Basin Float Array: the two float models", h1),
          Paragraph("Float programme, %s. Configuration as deployed; revised %s when the first model A floats "
                    "went in." % (INSTITUTE, rev.strftime("%B %Y")), small), Spacer(1, 6)]
    rows = [["", "Model K", "Model A"],
            ["Floats", "%s to %s" % (k[0], k[-1]), "%s to %s" % (a[0], a[-1])],
            ["Deployed", "2023 to 2024", "from 2025"],
            ["Telemetry", "Argos, one-way; the float stays at the surface for several hours while satellite passes "
                          "collect its messages", "Iridium, two-way; the profile is sent in minutes"],
            ["Position", "Argos fixes during the surface period", "GPS fix at the surface"],
            ["Cycle", "10 days", "10 days"],
            ["Parking pressure", "1,000 dbar", "1,000 dbar"],
            ["Profile pressure", "2,000 dbar", "2,000 dbar"],
            ["CTD", "pumped, conductivity cell with anti-fouling", "pumped, conductivity cell with anti-fouling"],
            ["Sampling on ascent", "5 dbar to 100 dbar, 10 dbar to 300, 25 dbar to 1,000, 50 dbar below",
             "as model K"]]
    t = Table([[Paragraph(c, cell) for c in r] for r in rows], colWidths=[34 * mm, 68 * mm, 68 * mm], repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA0AB")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E2E7F0")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    st += [t, Spacer(1, 8), Paragraph("The cycle", h2),
           Paragraph("Each cycle starts at the surface. The float sinks to its parking pressure in four to five hours, "
                     "drifts there for about nine days, sinks to 2,000 dbar and then rises to the surface at about 0.09 "
                     "metres a second, sampling as it goes, so the ascent takes about six hours. At the surface it sends "
                     "the profile and its technical message, then starts the next cycle.", body),
           Paragraph("Floats %s and %s are also set to sample on the way down to the parking pressure; that "
                     "descending profile is sent with the next ascending one." % (", ".join(down[:-1]), down[-1]), body),
           Paragraph("Profiles and real-time checks", h2),
           Paragraph("The data centre decodes each file as it arrives and runs the real-time checks, which flag each "
                     "value 1 (good), 2 (probably good), 3 (probably bad) or 4 (bad). A file that breaks off during "
                     "transmission is decoded as far as it goes and replaced when the float sends the profile again.",
                     body)]

    def foot(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(20 * mm, 10 * mm, "%s | Float programme | IBFA float models" % INSTITUTE)
        canvas.drawRightString(A4[0] - 20 * mm, 10 * mm, "Page %d" % d.page)
        canvas.restoreState()
    doc.build(st, onFirstPage=foot, onLaterPages=foot)
    postprocess.scrub(path, PEOPLE["eng"][0], rev.isoformat())


# ------------------------------------------------------------------ the data dictionary and readme (md)
def dictionary(F):
    return f"""# Data dictionary

## {F['register']}
One row per float of the Iberian Basin Float Array.

| Column | Meaning |
|---|---|
| wmo | the float's WMO number |
| serial | manufacturer's serial number |
| model | float model, K or A |
| telecom | telemetry, Argos or Iridium |
| ctd_serial | serial number of the float's CTD |
| deployed | deployment date |
| deploy_lat, deploy_lon | deployment position, decimal degrees |
| profiles | the profiles the float is set to record |
| cycle_days, park_dbar, profile_dbar | cycle length (days), parking pressure and profile pressure (dbar) |

## {F['profiles']}
One row per profile file the data centre decoded for the year: the array's floats and the partner float decoded beside them.

| Column | Meaning |
|---|---|
| profile_id | the data centre's id for the decoded file |
| wmo, cycle | float and cycle number |
| direction | A ascending, D descending |
| juld | the profile's date and time, days since 1950-01-01 00:00 UTC |
| latitude, longitude | the profile's position, decimal degrees |
| position_qc | 1 a fix, 8 estimated |
| file_version | 1 the first file received for the profile, 2 a later file for the same profile |

## {F['levels']}
One row per pressure level of a profile file.

| Column | Meaning |
|---|---|
| profile_id | the profile file |
| pres | sea pressure, dbar |
| temp | in-situ temperature, ITS-90, degrees C |
| psal | practical salinity |
| pres_qc, temp_qc, psal_qc | real-time flags: 1 good, 2 probably good, 3 probably bad, 4 bad |

## {F['technical']}
One row per float cycle: the cycle's timing and engineering values from the float's technical message. Times are UTC.

| Column | Meaning |
|---|---|
| wmo, cycle | float and cycle number |
| descent_start | start of the descent from the surface |
| park_start, park_end | start and end of the drift at parking pressure |
| ascent_start | start of the ascent from 2,000 dbar |
| ascent_end | end of the ascent |
| battery_v | battery voltage at the surface |
| vacuum_mbar | internal pressure |
| pump_s | buoyancy pump run time on the ascent, seconds |

## {F['deliveries']}
The data centre's receipt log for profile files of 2027 cycles, to the morning of 7 January 2028.

| Column | Meaning |
|---|---|
| received | when the data centre's server received the file, on the server's clock (Lisbon local time) |
| wmo, cycle, direction | the profile the file carries |
| file_name | file name as received |
| file_version | 1 the first file for the profile, 2 a later one |
| n_levels | levels decoded from the file |

## {F['arch_casts']} and {F['arch_levels']}
The institute's CTD archive: one row per cast, and one row per 10 dbar level of a cast.

| Column | Meaning |
|---|---|
| cast_id | cruise and station |
| cruise, station | cruise code and station number |
| time_utc | start of the cast, UTC |
| latitude, longitude | decimal degrees |
| temp_scale | temperature scale of the cast's temperatures, IPTS-68 or ITS-90 |
| max_pres | deepest level of the cast, dbar |
| pres | pressure, dbar |
| temp | in-situ temperature in the cast's scale, degrees C |
| psal | practical salinity |

## {F['intl_casts']} and {F['intl_levels']}
The RHD extract for the box: one row per station, and one row per level.

| Column | Meaning |
|---|---|
| station_id | the RHD's station id |
| platform, country | ship and country of the submitting institute |
| date | date of the cast, UTC |
| time | time of the cast, decimal hours UTC |
| lat, lon | decimal degrees, to 0.01 |
| max_depth_m | deepest level, metres |
| depth_m | depth, metres |
| temperature | in-situ temperature, ITS-90, degrees C |
| salinity | practical salinity |

## {F['recoveries']}
Floats of the programme's earlier arrays that were recovered, and what the laboratory found when it checked their sensors against a reference bath.

| Column | Meaning |
|---|---|
| wmo | the float |
| deployed, recovered | dates |
| reason | why the float was recovered |
| lab_salinity_error | salinity the float's CTD read less the bath's, practical salinity |
| lab_temperature_error | temperature the float's CTD read less the bath's, degrees C |

## {F['oxygen']}
Dissolved oxygen from the six model A floats carrying an optode, at every second level of the ascending profile.

| Column | Meaning |
|---|---|
| profile_id | the profile file |
| pres | pressure, dbar |
| doxy | dissolved oxygen, micromoles per kilogram |
| doxy_qc | real-time flag |
"""


def readme(F):
    return f"""# IBFA 2027 delayed-mode salinity release

The {INSTITUTE}'s Ocean Data Centre: the Iberian Basin Float Array's profiles for 2027, the reference data frozen for
the year's delayed-mode check, and the records behind the release report. Drawn on 7 January 2028.

| File | What it is |
|---|---|
| {F['thread']} | the release thread, 13 December 2027 to 7 January 2028 |
| {F['refnote']} | the reference data note |
| {F['models']} | the float programme's summary of the two float models |
| {F['register']} | the float register |
| {F['profiles']} | the year's decoded profile files |
| {F['levels']} | their levels |
| {F['technical']} | each cycle's timing and engineering values |
| {F['deliveries']} | the data centre's receipt log |
| {F['arch_casts']} | the institute's CTD archive, casts |
| {F['arch_levels']} | the institute's CTD archive, levels |
| {F['intl_casts']} | the RHD extract, stations |
| {F['intl_levels']} | the RHD extract, levels |
| {F['recoveries']} | laboratory checks of sensors from recovered floats |
| {F['oxygen']} | dissolved oxygen from the floats with optodes |
| DATA_DICTIONARY.md | the columns of every table |
"""
