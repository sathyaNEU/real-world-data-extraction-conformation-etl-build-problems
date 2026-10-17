"""task241 golden deliverables: the delayed-mode salinity workbook and the offsets chart."""
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Patch
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import params as P
import postprocess

XLSX = "ioa_dmqc_salinity_2027.xlsx"
PNG = "ioa_dmqc_offsets_2027.png"
NAVY, RUST, TEAL, INK = "1F3A5F", "B5482B", "2E7D7A", "1E2430"
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November",
          "December"]
WORDS = {4: "Four", 5: "Five", 6: "Six", 7: "Seven", 8: "Eight", 2: "two", 3: "three", 1: "one"}


def r3(x):
    """three decimals with no negative zero"""
    v = round(x, 3)
    return 0.0 if v == 0 else v


def _head(ws, title, sub):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13, color=INK)
    ws["A2"] = sub
    ws["A2"].font = Font(italic=True, size=9, color="555555")


def _table(ws, row, header, rows, widths, formats):
    thin = Side(style="thin", color="B9C0CC")
    for j, h in enumerate(header, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(bottom=thin)
    for i, r in enumerate(rows, row + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            if formats[j - 1]:
                c.number_format = formats[j - 1]
                c.alignment = Alignment(horizontal="right", vertical="top")
            else:
                c.alignment = Alignment(vertical="top")
            c.border = Border(bottom=thin)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def workbook(path, A):
    wb = Workbook()
    ws = wb.active
    ws.title = "Offsets 2027"
    adj = A["adjusted"]
    fresh = [w for w in adj if A["off"][w] < 0]
    salty = [w for w in adj if A["off"][w] > 0]
    _head(ws, "%s of the 24 floats have their 2027 salinity adjusted: %s reading fresh, %s reading salty"
          % (WORDS[len(adj)], WORDS[len(fresh)].lower(), WORDS[len(salty)]),
          "Offset: salinity reported less the reference water's salinity on the profile's potential-temperature levels "
          "at the profile's position, good values 1,400 to 1,900 dbar, averaged over the float's 2027 ascending "
          "profiles. Adjusted where more than %.3f in magnitude." % P.TOL)
    rows = []
    for w in sorted(A["off"]):
        rows.append([int(w), A["model"][w], A["used"][w], r3(A["off"][w]),
                     "Adjust" if w in adj else "Release as measured"])
    _table(ws, 4, ["WMO", "Model", "Ascending profiles", "Offset 2027 (practical salinity)", "Decision"], rows,
           [11, 8, 12, 16, 22], ["0", None, "0", "+0.000;-0.000;0.000", None])
    for i, row in enumerate(rows, 5):
        if row[4] == "Adjust":
            for j in range(1, 6):
                ws.cell(row=i, column=j).font = Font(bold=True, color=RUST)
    ws.cell(row=6 + len(rows), column=1, value="Source: %s, Ocean Data Centre; IBFA profiles received to 7 January "
            "2028, reference data frozen 15 December 2027." % P.INSTITUTE).font = Font(italic=True, size=9,
                                                                                       color="555555")
    w2 = wb.create_sheet("Reference casts")
    lo, hi = min(A["casts"].values()), max(A["casts"].values())
    _head(w2, "Reference casts reaching 2,000 dbar: %d to %d a year from 2015 to 2025" % (lo, hi),
          "Distinct casts in the institute's CTD archive and the RHD extract for the box, by year of the cast; the RHD "
          "copies of the institute's own casts counted once, RHD depths taken to pressure.")
    rows = [[y, A["casts_arch"][y], A["casts_other"][y], A["casts"][y]] for y in sorted(A["casts"])]
    rows.append(["2015 to 2025", sum(A["casts_arch"].values()), sum(A["casts_other"].values()),
                 sum(A["casts"].values())])
    _table(w2, 4, ["Year", "Institute archive", "Other institutes (RHD)", "Casts reaching 2,000 dbar"], rows,
           [13, 14, 16, 16], [None, "0", "0", "0"])
    for j in range(1, 5):
        w2.cell(row=4 + len(rows), column=j).font = Font(bold=True)
    w3 = wb.create_sheet("Delivery 2027")
    on, tot = sum(A["timely"].values()), sum(A["surfaced"].values())
    _head(w3, "%d of the array's %s ascending profiles in 2027 reached the data centre in full within 24 hours"
          % (on, "{:,}".format(tot)),
          "By the month the float surfaced (end of the ascent, from the cycle's technical message); a profile counts "
          "from the file that carried every level, the receipt log's Lisbon times taken to UTC.")
    rows = [[MONTHS[m - 1], A["surfaced"][m], A["timely"][m], A["timely"][m] / A["surfaced"][m]] for m in range(1, 13)]
    rows.append(["2027", tot, on, on / tot])
    _table(w3, 4, ["Month surfaced", "Ascending profiles", "In full within 24 hours", "Share"], rows,
           [15, 12, 14, 9], [None, "0", "0", "0.0%"])
    for j in range(1, 5):
        w3.cell(row=4 + len(rows), column=j).font = Font(bold=True)
    w4 = wb.create_sheet("Notes")
    notes = [
        "Release: the IBFA 2027 delayed-mode salinity release, due at the assembly centre on 21 January 2028. Tolerance "
        "and layer as set in the release thread of 13 December 2027.",
        "Float data: each float's 2027 ascending profiles, one file per profile (where a profile was received twice, the "
        "file that carried every level), levels at flag 1 or 2 in pressure, temperature and salinity, 1,400 to 1,900 "
        "dbar. The partner institute's float decoded beside the array is not included.",
        "Reference water: potential temperature (referred to the surface) computed for every float level and every "
        "reference level; the reference casts' salinity at 1,300 to 2,050 dbar fitted on potential temperature with "
        "terms for longitude and latitude, and read at each float level's potential temperature and the profile's "
        "position. The 2015 and 2016 cruise temperatures were taken from IPTS-68 to ITS-90 and the RHD depths to "
        "pressure first.",
        "The deep layer sat deeper in 2027 than in the reference years and moved with eddies, and on a given "
        "potential-temperature level the water freshens westward across the basin; the offset is therefore read on "
        "temperature levels at the profile's own position, not at fixed pressure or against the basin's average.",
        "Reference casts: a cast in the RHD extract is the institute's own when its time is within 15 minutes and its "
        "position within 0.02 degree of an archive cast; repeat casts at a station a few hours apart are separate "
        "casts. RHD depths are taken to pressure (Saunders 1981) before the 2,000 dbar test.",
        "Delivery: surfacing is the end of the ascent in the cycle technical file. The model K floats time-stamp their "
        "profiles at the start of the ascent, about six hours before they surface, so the profile time is not used. The "
        "receipt log is on the server's Lisbon clock and is taken to UTC.",
        "Prepared by Naiara Nogueira, Ocean Data Centre, 7 January 2028.",
    ]
    w4["A1"] = "Notes"
    w4["A1"].font = Font(bold=True, size=12)
    for i, t in enumerate(notes, 3):
        w4.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True, vertical="top")
    w4.column_dimensions["A"].width = 112
    for sh in (ws, w2, w3):
        sh.sheet_view.showGridLines = False
        sh.print_title_rows = "4:4"
        sh.page_setup.orientation = "portrait"
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
        sh.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)
    postprocess.scrub(path, "Naiara Nogueira", "2028-01-07")
    postprocess.finish_office(path, "2028-01-07T09:40:00Z", "2028-01-07T16:55:00Z", P.INSTITUTE)


def chart(path, A):
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9})
    ws = sorted(A["off"])
    vals = [A["off"][w] for w in ws]
    fig, ax = plt.subplots(figsize=(10.6, 5.6), dpi=150)
    ax.axhspan(-P.TOL, P.TOL, color="#E3ECF3", zorder=0)
    for y in (-P.TOL, P.TOL):
        ax.axhline(y, color="#5C6B7A", lw=0.9, ls="--", zorder=1)
    ax.axhline(0, color="#333333", lw=0.6, zorder=1)
    cols = ["#" + RUST if w in A["adjusted"] else "#8FA3B8" for w in ws]
    ax.bar(range(len(ws)), vals, color=cols, width=0.7, zorder=2)
    for i, (w, v) in enumerate(zip(ws, vals)):
        if w in A["adjusted"]:
            ax.text(i, v + (0.0012 if v > 0 else -0.0012), "%+.3f" % r3(v), ha="center",
                    va="bottom" if v > 0 else "top", fontsize=8, color="#" + RUST, fontweight="bold")
    ax.set_xticks(range(len(ws)))
    ax.set_xticklabels(ws, rotation=60, ha="right", fontsize=8)
    ax.set_xlim(-0.7, len(ws) - 0.3)
    ax.set_ylim(-0.036, 0.030)
    ax.set_yticks([-0.03, -0.02, -0.01, 0.0, 0.01, 0.02, 0.03])
    ax.set_yticklabels(["%+.2f" % t if t else "0" for t in [-0.03, -0.02, -0.01, 0.0, 0.01, 0.02, 0.03]])
    ax.set_ylabel("2027 salinity offset (practical salinity)")
    ax.set_xlabel("Float (WMO number)")
    ax.text(len(ws) - 0.4, P.TOL + 0.0007, "tolerance +%.3f" % P.TOL, ha="right", va="bottom", fontsize=8,
            color="#5C6B7A")
    ax.text(len(ws) - 0.4, -P.TOL - 0.0007, "tolerance -%.3f" % P.TOL, ha="right", va="top", fontsize=8,
            color="#5C6B7A")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    adj = A["adjusted"]
    ax.set_title("%s IBFA floats fall outside the salinity tolerance in 2027 and are adjusted in the release"
                 % WORDS[len(adj)], fontsize=11.5, loc="left", fontweight="bold", color="#" + INK)
    ax.legend(handles=[Patch(color="#" + RUST, label="adjusted (%d)" % len(adj)),
                       Patch(color="#8FA3B8", label="released as measured (%d)" % (len(ws) - len(adj))),
                       Patch(color="#E3ECF3", label="tolerance band")], loc="lower left", frameon=False, fontsize=8)
    fig.text(0.01, 0.008, "Offset: salinity reported less the reference water's on potential-temperature levels at the "
             "profile's position, 1,400 to 1,900 dbar,\nmean of the float's 2027 ascending profiles. %s, Ocean Data "
             "Centre, 7 January 2028." % P.INSTITUTE, fontsize=7, color="#555555", linespacing=1.4)
    fig.tight_layout(rect=(0, 0.045, 1, 1))
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)


def write(folder, A):
    os.makedirs(folder, exist_ok=True)
    workbook(os.path.join(folder, XLSX), A)
    chart(os.path.join(folder, PNG), A)
