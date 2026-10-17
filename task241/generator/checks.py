"""task241 checks: the three rules executed on the simulator's clean truth (R1 to R3c), every semantic fork row (SF1 to
SF12), the rungs and the corridor of position-mapped methods, every Layer 1 slip against the call, the ask devices and
their stops, the separation of the call from the devices, the documents' pins, the input gates and the realism sweep."""
import datetime as dt
import itertools
import os
import re

import numpy as np
import pandas as pd

import estim as E
import params as P
import records as RC

N = [0]
GOLD = ["1990001", "1990004", "1990005", "1990007", "1990008", "1990012"]


def ok(cond, label):
    assert cond, label
    N[0] += 1


def _pdf_text(path):
    from pypdf import PdfReader
    return " ".join(" ".join(pg.extract_text().split()) for pg in PdfReader(path).pages)


def _docx_text(path):
    import docx
    d = docx.Document(path)
    return " ".join([p.text for p in d.paragraphs] + [c.text for t in d.tables for row in t.rows for c in row.cells])


# ------------------------------------------------------------------ the routes on the clean truth
def truth_frames(W, R):
    """the world before any record is written: each array float's 2027 levels in the layer (measured pressure and
    salinity, the water's true potential temperature) and every reference cast's levels (true pressure, potential
    temperature and salinity)"""
    rows = []
    for f in R.array:
        for c in f["cycles"]:
            m = (c["p"] >= P.LAYER[0]) & (c["p"] <= P.LAYER[1])
            for p, th, s in zip(c["p"][m], c["theta_true"][m], c["s"][m]):
                rows.append((f["wmo"], c["cycle"], c["lon"], c["lat"], p, th, s, c["delta"]))
    fl = pd.DataFrame(rows, columns=["wmo", "cycle", "lon", "lat", "pres", "theta", "psal", "delta"])
    refs = []
    for c in R.casts:
        m = (c["p"] <= c["pmax"]) & (c["p"] >= 1250) & (c["p"] <= 2050)
        for p, th, s in zip(c["p"][m], c["theta"][m], c["s"][m]):
            refs.append((c["id"], c["lon"], c["lat"], p, th, s))
    for c in R.others:
        m = (c["p"] >= 1250) & (c["p"] <= 2050)
        for p, th, s in zip(c["p"][m], c["theta"][m], c["s"][m]):
            refs.append((c["id"], c["lon"], c["lat"], p, th, s))
    ref = pd.DataFrame(refs, columns=["cast", "lon", "lat", "pres", "theta", "psal"])
    return fl, ref


def truth_route(fl, ref, route):
    fl = fl.copy()
    if route in ("golden", "theta_global", "self"):
        r = ref[(ref["pres"] >= 1300) & (ref["pres"] <= 2050)]
        beta = E.ref_model(r, position=(route != "theta_global"))
        fl["d"] = fl["psal"] - E.ref_s(beta, fl["theta"].values, fl["lon"].values, fl["lat"].values)
    else:
        grid = np.arange(1305.0, 1996.0, 10.0)
        prof = {}
        for c, g in ref.groupby("cast"):
            g = g.sort_values("pres")
            if g["pres"].min() <= grid[0] and g["pres"].max() >= grid[-1]:
                prof[c] = (np.interp(grid, g["pres"].values, g["psal"].values), g["lon"].iloc[0], g["lat"].iloc[0])
        names = list(prof)
        M = np.array([prof[c][0] for c in names])
        clon = np.array([prof[c][1] for c in names])
        clat = np.array([prof[c][2] for c in names])
        d = np.zeros(len(fl))
        for (w, cyc), idx in fl.groupby(["wmo", "cycle"]).groups.items():
            g = fl.loc[idx]
            if route == "p_global":
                refp = M.mean(axis=0)
            else:
                dist = ((clon - g["lon"].iloc[0]) * np.cos(np.radians(g["lat"].iloc[0]))) ** 2 + \
                    (clat - g["lat"].iloc[0]) ** 2
                refp = M[np.argsort(dist)[:12]].mean(axis=0)
            d[fl.index.get_indexer(idx)] = g["psal"].values - np.interp(g["pres"].values, grid, refp)
        fl["d"] = d
    per = fl.groupby(["wmo", "cycle"])["d"].mean().reset_index().sort_values(["wmo", "cycle"])
    out = {}
    for w, g in per.groupby("wmo"):
        out[w] = float(g["d"].mean() - (g["d"].iloc[0] if route == "self" else 0.0))
    return out


def run(target, golden, o, X, A):
    W, R = o["W"], o["R"]
    F = X["files"]
    D = X["D"]
    off, adj, routes = X["off"], X["adjusted"], X["routes"]
    ok(adj == GOLD, "golden: six floats adjusted (%s)" % adj)

    # ---------------------------------------------------------------- the truth behind the golden
    true = {f["wmo"]: float(np.mean([c["delta"] for c in f["cycles"]])) for f in R.array}
    ok(sorted(w for w, v in true.items() if abs(v) > P.TOL) == GOLD, "truth: the six floats' sensors are off by more "
       "than the tolerance over 2027")
    ok(max(abs(off[w] - true[w]) for w in true) < 0.0012, "the golden offsets reproduce each float's true 2027 "
       "sensor error to within 0.0012 (max %.4f)" % max(abs(off[w] - true[w]) for w in true))
    ok(min(abs(true[w]) for w in GOLD) > 0.0195 and max(abs(true[w]) for w in true if w not in GOLD) < 0.0035,
       "truth margins: adjusted floats 0.0195 or more, the rest under 0.0035")
    kinds = {f["wmo"]: f["kind"] for f in R.array}
    ok(sorted(w for w in GOLD if kinds[w] == "drift") == A["fresh"] and
       sorted(w for w in GOLD if kinds[w] == "offset") == A["salty"], "the fresh four drifted, the salty two were off "
       "from deployment")

    # ---------------------------------------------------------------- R1: the field's raw material ships
    for f in os.listdir(target):
        if f.endswith(".csv"):
            cols = pd.read_csv(os.path.join(target, f), nrows=1).columns
            ok(not any(re.search(r"offset|adjust|drift|delta|theta|potential|error_est|reference_s|anomal", c, re.I)
                       for c in cols if c != "lab_salinity_error" and c != "lab_temperature_error"),
               "R1 no shipped column holds an offset, a potential temperature or a reference salinity (%s)" % f)
    ok(len(D["lev"]) > 25000 and {"pres", "temp", "psal"} <= set(D["lev"].columns) and
       {"pres", "temp", "psal"} <= set(D["al"].columns) and {"depth_m", "temperature", "salinity"} <= set(D["il"].columns),
       "R1 the float levels and the reference casts' levels ship as measured (pressure or depth, in-situ "
       "temperature, practical salinity)")

    # ---------------------------------------------------------------- R3: the analyst test on the clean truth
    fl, ref = truth_frames(W, R)
    tr = {r: truth_route(fl, ref, r) for r in ("p_global", "p_local", "theta_global", "golden", "self")}
    a0 = E.adjusted(tr["p_global"])
    ok(a0 != GOLD and len(set(a0) - set(GOLD)) >= 6, "R3a the default path (every good value at fixed pressure "
       "against the reference casts' mean) on the clean truth adjusts %d floats, %d of them clean" %
       (len(a0), len(set(a0) - set(GOLD))))
    ok(E.adjusted(tr["golden"]) == GOLD, "R3b the decisive step alone (potential-temperature levels, reference at the "
       "profile's position) on the clean truth gives the six")
    ok(E.adjusted(tr["p_local"]) != GOLD, "R3b position alone at fixed pressure stays wrong on the clean truth (%d)"
       % len(E.adjusted(tr["p_local"])))
    ok(E.adjusted(tr["theta_global"]) != GOLD, "R3b temperature levels alone (the pooled reference) stay wrong on the "
       "clean truth (%d)" % len(E.adjusted(tr["theta_global"])))
    ok(E.adjusted(tr["self"]) == sorted(w for w in GOLD if kinds[w] == "drift"), "R3b drift from the first profile "
       "finds only the drifting floats on the clean truth")
    # the default path on the shipped files, every filed rule applied
    ok(len(E.adjusted(routes["p_global"])) == 14 and set(GOLD) <= set(E.adjusted(routes["p_global"])),
       "default path on the pack ends on 14 floats, all six among them, eight clean")
    ok(len(E.adjusted(routes["theta_global"])) == 13 and set(GOLD) <= set(E.adjusted(routes["theta_global"])),
       "pooled temperature-level rung on the pack: 13")
    ok(len(E.adjusted(routes["p_local"])) == 13, "nearest-casts rung at fixed pressure on the pack: 13")
    ok(E.adjusted(routes["self"]) == A["fresh"], "drift-only rung on the pack: the four drifting floats")
    for n, v in A["p_variants"].items():
        ok(13 <= v["n"] <= 17 and v["extra"] >= 7, "fixed-pressure variant %s adjusts %d, %d clean" % (n, v["n"], v["extra"]))
    for n, v in A["t_variants"].items():
        ok(10 <= v["n"] <= 14 and v["extra"] >= 4 and not v["missed"], "pooled temperature-level variant %s adjusts %d, "
           "%d clean" % (n, v["n"], v["extra"]))
    ok(tuple(A["p_range"]) == (13, 17) and A["p_extra_min"] == 7 and tuple(A["t_range"]) == (10, 14) and
       A["t_extra_min"] == 4, "the write-up's rung ranges are the measured ones")
    ok(A["theta_span"][0] < A["theta_span"][2] and A["theta_span"][3] < A["theta_span"][1], "the reference levels' "
       "potential temperatures span every float level's used")
    # why the default path fails: the layer sat deeper in 2027, and the water freshens westward on a level
    pr, lv = E.float_layer(D)
    r = E.reference(D)
    def iso(g, th0=5.5):
        g = g.sort_values("pres")
        t, p = g["theta"].values, g["pres"].values
        return float(np.interp(th0, t[::-1], p[::-1])) if t.min() <= th0 <= t.max() else np.nan
    fz = np.nanmean([iso(g) for _, g in lv.groupby("profile_id")])
    rz = np.nanmean([iso(g) for _, g in r.groupby("cast")])
    ok(fz - rz > 10.0, "the 5.5 degree isotherm sat %.1f dbar deeper in the 2027 float profiles than in the reference "
       "casts" % (fz - rz))
    beta = E.ref_model(r)
    ok(beta[3] * 2.0 > 0.0025, "on the 5.5 degree level the reference water freshens westward, %.4f per degree of "
       "longitude" % (beta[3] * 2.0))

    # ---------------------------------------------------------------- R3c: no plant
    for k in ("thread", "refnote", "models", "readme", "dict"):
        p = os.path.join(target, F[k])
        txt = _pdf_text(p) if p.endswith(".pdf") else _docx_text(p) if p.endswith(".docx") else open(p).read()
        ok(not re.search(r"0\.0(?!10\b)[1-4]\d|offset of|drift(ed|ing) float|potential temperature|isotherm|heave|"
                         r"at the (profile|float)'s (own )?position|mapped|fixed pressure", txt, re.I),
           "R3c no document carries an offset, a drift verdict or the method (%s)" % k)
    th = open(os.path.join(target, F["thread"]), encoding="utf-8").read()
    ok(not re.search(r"potential|isotherm|heave|theta|mapped|position|pressure level|fixed pressure", th, re.I),
       "R3c the release thread files the estimand and names no method")
    ok("true salinity of the water it sampled" in th and "more than 0.010 in magnitude" in th and
       "between 1,400 and 1,900 dbar" in th and "flag 1 or 2" in th, "the estimand, tolerance, layer and good values "
       "are filed in the release thread")
    ok("in full within 24 hours of the float surfacing" in th and "by the month the float surfaced" in th,
       "the delivery record is filed in the release thread")

    # ---------------------------------------------------------------- the corridor of position-mapped methods (C1, C3)
    for name, kw in [("near5", dict(k=5)), ("near10", dict(k=10)), ("near20", dict(k=20)), ("near60", dict(k=60)),
                     ("near100", dict(k=100)), ("rad1", dict(radius=1.0)), ("rad2", dict(radius=2.0)),
                     ("rad3", dict(radius=3.0)), ("rad4", dict(radius=4.0)), ("gauss2", dict(scale=2.0)),
                     ("gauss3", dict(scale=3.0)), ("gauss4", dict(scale=4.0)), ("insitu10", dict(k=10, insitu=True)),
                     ("ext10", dict(k=10, sources=("extract",))), ("arch10", dict(k=10, sources=("archive",)))]:
        v = E.theta_local(D, **kw)
        ok(E.adjusted(v) == GOLD, "C1 %s on potential-temperature levels at the profile's position gives the six" % name)
        if name in ("near5", "near10", "near20", "rad1", "rad2", "gauss2", "insitu10", "ext10"):
            ok(max(abs(v[w] - off[w]) for w in off) <= 0.003, "C1 %s reproduces every offset within 0.003 (max %.4f)"
               % (name, max(abs(v[w] - off[w]) for w in off)))
    v = E.theta_local(D, radius=5.0)
    ok(E.adjusted(v) == GOLD, "C3 the corridor reaches a 5 degree radius")

    # ---------------------------------------------------------------- semantic fork rows
    reg = set(D["reg"]["wmo"])
    tech = D["tech"].copy()
    m = D["prof"][D["prof"]["direction"] == "A"].merge(tech[["wmo", "cycle", "ascent_end"]], on=["wmo", "cycle"])
    jt = pd.to_datetime("1950-01-01") + pd.to_timedelta((m["juld"] * 86400).round(), unit="s")
    ok(int((jt.dt.year != pd.to_datetime(m["ascent_end"]).dt.year).sum()) == 0,
       "SF1 2027 profiles: the year of juld and the year of surfacing select the same profiles")
    lvall = D["lev"].merge(pr[["profile_id"]], on="profile_id")
    lay = lvall[(lvall["pres"] >= P.LAYER[0]) & (lvall["pres"] <= P.LAYER[1])]
    Dp = dict(D, lev=D["lev"].assign(pres_qc=1, temp_qc=1))
    ok(E.adjusted(E.offsets(Dp, "golden")) == GOLD, "SF2 good values: salinity's flag alone (%d pressure and %d "
       "temperature values flagged in the layer) gives the same call" % (int((~lay["pres_qc"].isin([1, 2])).sum()),
                                                                     int((~lay["temp_qc"].isin([1, 2])).sum())))
    ok(int((D["lev"]["temp_qc"] == 4).sum()) > 0 and int((D["lev"]["pres_qc"] == 4).sum()) > 0,
       "the real-time checks flag temperature and pressure as well as salinity")
    ok(E.adjusted(E.offsets(D, "golden", dedup=None)) == GOLD, "SF3 a profile received twice, both files kept: same call")
    D2 = dict(D, reg=pd.concat([D["reg"], pd.DataFrame({"wmo": A["partner"]})], ignore_index=True))
    o2 = E.offsets(D2, "golden")
    ok(E.adjusted(o2) == GOLD and len(o2) == 25, "SF4 the partner float included: it is not adjusted, same call")
    ok(min(abs(off[w]) for w in GOLD) - P.TOL > 0.009 and P.TOL - max(abs(off[w]) for w in off if w not in GOLD) > 0.006,
       "SF5 more than 0.010 or at least 0.010: no offset near the tolerance")
    lv2 = lv.copy()
    rb = E.reference(D)
    beta = E.ref_model(rb)
    lv2["d"] = lv2["psal"] - E.ref_s(beta, lv2["theta"].values, lv2["longitude"].values, lv2["latitude"].values)
    pooled = lv2.groupby("wmo")["d"].mean().to_dict()
    ok(E.adjusted(pooled) == GOLD and max(abs(pooled[w] - off[w]) for w in off) < 0.0005,
       "SF6 the mean over all the float's layer values against the mean of profile means: same call")
    ok(E.adjusted({w: v * 35.16504 / 35.0 for w, v in off.items()}) == GOLD, "SF7 absolute salinity against practical "
       "salinity: same call")
    ok(E.adjusted(E.theta_local(D, k=10, insitu=True)) == GOLD, "SF8 in-situ temperature levels: same call")
    ok(E.adjusted(E.theta_local(D, k=10, sources=("archive",))) == GOLD and
       E.adjusted(E.theta_local(D, k=10, sources=("extract",))) == GOLD, "SF9 either reference source alone: same call")
    ok(E.adjusted(E.offsets(D, "golden", layer=(1400.5, 1899.5))) == GOLD, "SF10 the layer's ends inclusive or "
       "exclusive: same call")
    ok(all(np.sign(off[w]) == np.sign(true[w]) for w in GOLD), "SF11 the offset's sign: float less reference, the fresh "
       "floats negative")
    ok(E.adjusted({w: abs(v) for w, v in off.items()}) == GOLD, "SF12 more than 0.010 in magnitude, either sign")

    # ---------------------------------------------------------------- Layer 1 slips leave the call
    ok(E.adjusted(E.offsets(D, "golden", good=False)) == GOLD, "L1 flagged values kept: same call")
    ok(E.adjusted(E.offsets(D, "golden", ascending=False)) == GOLD, "L1 descending profiles kept: same call")
    Dn = dict(D, ac=D["ac"].assign(temp_scale="ITS-90"))
    ok(E.adjusted(E.offsets(Dn, "golden")) == GOLD, "L1 IPTS-68 read as ITS-90: same call")
    Dd = dict(D)
    il = D["il"].copy()
    ic = D["ic"]
    Dd["il"] = il.assign(depth_m=E.pres_from_depth(il["depth_m"].values, il.merge(ic[["station_id", "lat"]],
                                                                                 on="station_id")["lat"].values))
    ok(E.adjusted(E.offsets(Dd, "golden")) == GOLD, "L1 RHD depth read as pressure: same call")
    keep = []
    acw = D["ac"].assign(w=pd.to_datetime(D["ac"]["time_utc"]))
    icw = ic.assign(w=pd.to_datetime(ic["date"]) + pd.to_timedelta((ic["time"] * 3600).round(), unit="s"))
    for _, row in icw.iterrows():
        mm = ((acw["w"] - row["w"]).abs() <= pd.Timedelta(minutes=15)) & ((acw["latitude"] - row["lat"]).abs() <= 0.02) \
            & ((acw["longitude"] - row["lon"]).abs() <= 0.02)
        keep.append(not mm.any())
    Du = dict(D, ic=ic[keep], il=D["il"][D["il"]["station_id"].isin(set(ic[keep]["station_id"]))])
    ou = E.offsets(Du, "golden")
    ok(E.adjusted(ou) == GOLD and max(abs(ou[w] - off[w]) for w in off) < 0.001, "L1 the RHD copies of the institute's "
       "casts dropped from the reference: same call, offsets within 0.001")

    # ---------------------------------------------------------------- separation: the ask devices sit outside the call
    ok(all(w in reg for w in pr["wmo"].unique()), "separation: the call's profiles are the array's ascending profiles")
    ok(not set(A["partner"]) & set(pr["wmo"]), "separation: the partner float is outside the call's population")

    # ---------------------------------------------------------------- ask A: reference casts reaching 2,000 dbar
    truthA = {y: 0 for y in range(2015, 2026)}
    for c in R.casts:
        truthA[c["year"]] += int(c["pmax"] >= 2000)
    for c in R.others:
        truthA[c["year"]] += int(c["pmax"] >= 2000)
    ok(X["casts"] == truthA, "ask A: the golden counts are the distinct casts reaching 2,000 dbar each year")
    ok(sum(1 for c in R.casts if c.get("dup")) == A["matched"], "ask A: every RHD copy of an institute cast is matched "
       "(%d), none wrongly" % A["matched"])
    stops = {"natural": E.ref_casts_2000(D, dedup=None, convert=False),
             "no_dedup": E.ref_casts_2000(D, dedup=None),
             "depth_as_pressure": E.ref_casts_2000(D, convert=False),
             "loose_dedup": E.ref_casts_2000(D, dedup="loose")}
    for k, v in stops.items():
        moved = sum(v[y] != truthA[y] for y in truthA)
        need = {"natural": 11, "no_dedup": 11, "depth_as_pressure": 8, "loose_dedup": 7}[k]
        ok(moved >= need, "ask A stop %s moves %d of 11 years" % (k, moved))
    ok(all(stops["no_dedup"][y] > truthA[y] for y in truthA), "ask A primary: the RHD copies counted again raise every "
       "year")
    ick = D["ic"].copy()
    pm = E.pres_from_depth(ick["max_depth_m"].values, ick["lat"].values)
    ok(np.min(np.abs(pm - 2000.0)) > 3.0, "ask A: no RHD station within 3 dbar of 2,000 (closest %.1f)" %
       np.min(np.abs(pm - 2000.0)))
    ok(np.min(np.abs(D["ac"]["max_pres"] - 2000.0)) > 3.0, "ask A: no archive cast within 3 dbar of 2,000")
    md = D["ic"]["max_depth_m"].values
    ok(not ((md > 1962) & (md < 1981)).any(), "ask A: no RHD station's deepest level between 1,962 and 1,981 m, where "
       "the choice of depth-to-pressure conversion would decide it")
    conv = {"saunders": E.pres_from_depth(md, D["ic"]["lat"].values), "factor_1.01": md * 1.01,
            "factor_1.0094": md * 1.0094 + 2.2e-6 * md ** 2}
    for k, pm2 in conv.items():
        ok(np.array_equal(pm2 >= 2000, conv["saunders"] >= 2000), "ask A: the %s conversion selects the same RHD "
           "stations as reaching 2,000 dbar" % k)
    ok(len(D["ac"]["max_pres"].unique()) >= 8, "the archive's deepest levels differ from cast to cast (%d values)"
       % len(D["ac"]["max_pres"].unique()))
    plat = D["ic"].groupby("platform")["country"].nunique()
    ok((plat == 1).all(), "each RHD platform has one country")
    # the dedup rule's margins: copies within 3 minutes and 0.005 degree, repeats 3 hours or more apart
    reps = [c for c in R.casts if c.get("repeat")]
    ok(len(reps) >= 15, "ask A hazard: %d repeat casts at a station a few hours apart" % len(reps))

    # ---------------------------------------------------------------- ask B: deliveries in full within 24 hours
    cut = dt.datetime.strptime(RC.CUTOFF, "%Y-%m-%dT%H:%M:%SZ")
    truthB = {mth: 0 for mth in range(1, 13)}
    for f in R.array:
        for c in f["cycles"]:
            truthB[c["time"].month] += int(c["delivered"] <= cut and (c["delivered"] - c["time"]) <=
                                           dt.timedelta(hours=24))
    ok(X["timely"] == truthB, "ask B: the golden counts are the profiles whose complete file arrived within 24 hours of "
       "surfacing")
    ok(E.timely_by_month(D, repeated="winter") == X["timely"], "ask B: the hour the clocks went back read either way "
       "gives the same counts")
    dl = D["dlv"].copy()
    dl["at"] = E.lisbon_to_utc(dl["received"])
    dl = dl[dl["direction"] == "A"].sort_values("n_levels").drop_duplicates(["wmo", "cycle"], keep="last")
    x = m.merge(dl[["wmo", "cycle", "at"]], on=["wmo", "cycle"])
    x = x[x["wmo"].isin(reg)]
    sec = (x["at"] - pd.to_datetime(x["ascent_end"]).dt.tz_localize(None)).dt.total_seconds() - 86400
    ok(sec[sec > 0].min() >= 30 and -sec[sec <= 0].max() >= 30, "ask B: no profile within 30 seconds of the 24 hour "
       "line on either side (closest %.0f s late, %.0f s early); both clocks are kept to the second" %
       (sec[sec > 0].min(), -sec[sec <= 0].max()))
    for combo in itertools.product(["ascent_end", "juld"], ["complete", "first"], ["lisbon", "utc"], [False, True],
                                   [False, True]):
        if combo == ("ascent_end", "complete", "lisbon", False, False):
            continue
        v = E.timely_by_month(D, anchor=combo[0], version=combo[1], clock=combo[2], partners=combo[3],
                              descending=combo[4])
        moved = sum(v.get(mth, 0) != truthB[mth] for mth in truthB)
        need = 12 if combo[0] == "juld" and not (combo[1] == "first" and combo[3]) else 1
        ok(moved >= need, "ask B stop %s moves %d of 12 months" % (combo, moved))
    ok(all(E.timely_by_month(D, anchor="juld")[mth] < truthB[mth] for mth in truthB),
       "ask B primary: the profile time read as surfacing lowers every month")

    # ---------------------------------------------------------------- devices outside the call's rows
    ok(E.adjusted(E.offsets(D, "golden")) == GOLD and set(A["partner"]).isdisjoint(set(pr["wmo"])),
       "separation: no device row in the call's population")

    # ---------------------------------------------------------------- input gates and realism
    files = os.listdir(target)
    ok(len(files) >= 10 and len({f.rsplit(".", 1)[1] for f in files}) >= 3, "input gates: %d files" % len(files))
    big = max(sum(1 for _ in open(os.path.join(target, f))) for f in files if f.endswith(".csv"))
    ok(big >= 25000, "input gates: a file of %d rows" % big)
    for f in files:
        p = os.path.join(target, f)
        if f.endswith((".csv", ".md", ".eml")):
            ok(not re.search("[\u2013\u2014]", open(p, encoding="utf-8").read()), "no en or em dash in %s" % f)
    reg_df = D["reg"]
    ok(reg_df["wmo"].is_unique and len(reg_df) == 24, "the register holds the 24 floats once each")
    ok(D["prof"]["profile_id"].is_unique and D["ac"]["cast_id"].is_unique and D["ic"]["station_id"].is_unique,
       "every shipped id is unique within its file")
    oxy = pd.read_csv(os.path.join(target, F["oxygen"]))
    ok(oxy["profile_id"].isin(D["prof"]["profile_id"]).all(), "the oxygen file's profiles are in the profile file")
    nopt = D["prof"][D["prof"]["profile_id"].isin(oxy["profile_id"])]["wmo"].nunique()
    ok(nopt == 6, "six floats carry an optode (%d)" % nopt)
    ok(all(D["reg"].set_index("wmo").loc[w, "model"] == "A" for w in
           D["prof"][D["prof"]["profile_id"].isin(oxy["profile_id"])]["wmo"].unique()), "the optode floats are model A")
    return N[0]
