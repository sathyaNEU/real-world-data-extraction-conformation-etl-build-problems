"""The two deliverables: the authority's return workbook and the chart that carries the decision."""
import os

import matplotlib
matplotlib.use("Agg")
matplotlib.rcParams["svg.hashsalt"] = "task189"
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import world as W

THIN = Side(style="thin", color="8A8A8A")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="E4E7DF")


def workbook(path, g, rg_names):
    wb = Workbook()
    ws = wb.active
    ws.title = "Return 2025-26"
    ws["A1"] = "%s" % W.AUTHORITY
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Annual return to the %s under regulation 7.6, reporting year 2025/26, "
                "1 April 2025 to 31 March 2026" % W.REGULATOR)
    ws["A3"] = "Listed substance: MCPA. Standard 0.05 ug/l, annual average (Schedule 2)."
    ws["A4"] = "Return due 30 November 2026. Prepared from the laboratory export of 16 November 2026."
    for r in (2, 3, 4):
        ws.cell(row=r, column=1).font = Font(size=9)
    hdr = ["Water body code", "Water body", "Compliance samples in the year",
           "Results below the reporting limit", "Assessment"]
    for i, h in enumerate(hdr, start=1):
        c = ws.cell(row=6, column=i, value=h)
        c.font = Font(bold=True)
        c.fill = HEAD
        c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical="bottom", horizontal="center")
    row = 7
    for code, o in sorted(g["assess"].items(), key=lambda kv: kv[1]["name"]):
        vals = [code, o["name"], o["n"], o["nd"],
                "non-compliant" if o["verdict"] == "non-compliant" else "compliant"]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=row, column=i, value=v)
            c.border = BOX
            if i >= 3:
                c.alignment = Alignment(horizontal="center")
            if i == 5 and v == "non-compliant":
                c.font = Font(bold=True)
        row += 1
    ws.cell(row=row + 1, column=1, value="Water bodies assessed as non-compliant for MCPA").font = \
        Font(bold=True)
    c = ws.cell(row=row + 1, column=3, value=g["count"])
    c.font = Font(bold=True)
    c.alignment = Alignment(horizontal="center")
    ws.cell(row=row + 2, column=1,
            value="Non-compliant water bodies: " + ", ".join(g["non_compliant"])).font = Font(size=9)
    ws.cell(row=row + 4, column=1, value=(
        "Three quarters of the MCPA results in the year were reported below a reporting limit, and "
        "the two methods in the programme report to 0.005 ug/l and 0.10 ug/l. The averages behind "
        "this return were estimated from the results as reported, each at its own reporting limit, "
        "and no value was substituted for a result below a limit.")).font = Font(size=9, italic=True)
    ws.cell(row=row + 5, column=1, value="Prepared by %s, catchment monitoring"
            % W.PEOPLE["monitoring_lead"]).font = Font(size=9)
    for col, wd in zip("ABCDE", [20, 22, 17, 17, 16]):
        ws.column_dimensions[col].width = wd
    ws.row_dimensions[6].height = 42
    wb.save(path)
    return path


def chart(path, g):
    items = sorted(g["assess"].values(), key=lambda o: -(o["n"]))
    names = [o["name"] for o in items]
    quant = [o["n"] - o["nd"] for o in items]
    below = [o["nd"] for o in items]
    nc = [o["verdict"] == "non-compliant" for o in items]
    fig, ax = plt.subplots(figsize=(10.2, 5.4), dpi=140)
    x = range(len(names))
    ax.bar(x, quant, color="#3d6b57", label="quantified results")
    ax.bar(x, below, bottom=quant, color="#cbd3c5", edgecolor="#8d9a86", linewidth=0.5,
           label="results below the reporting limit")
    for i, (o, flag) in enumerate(zip(items, nc)):
        ax.text(i, o["n"] + 0.9, "non-compliant" if flag else "", ha="center", va="bottom",
                fontsize=7.6, color="#8c2f1d", rotation=90)
    ax.set_xticks(list(x))
    ax.set_xticklabels(names, rotation=38, ha="right", fontsize=8.4)
    ax.set_ylabel("MCPA compliance samples in 2025/26", fontsize=9)
    ax.set_title("MCPA compliance samples by water body, Araglen River Basin District, 2025/26\n"
                 "%d of %d water bodies assessed as non-compliant with the 0.05 ug/l annual average"
                 % (g["count"], len(items)), fontsize=10.5, loc="left")
    ax.legend(frameon=False, fontsize=8.4, loc="upper right")
    ax.spines[["top", "right"]].set_visible(False)
    ax.set_ylim(0, max(o["n"] for o in items) * 1.28)
    ax.grid(axis="y", color="#e6e6e6", linewidth=0.6)
    ax.set_axisbelow(True)
    fig.tight_layout()
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)
    return path
