"""Write the bundle the authority would hand over: the laboratory export and the reference files."""
import csv
import datetime as dt
import json
import os
import random

import records as R
import sim
import world as W

FILES = {
    "export": "lims_results_export_2025_2026.csv",
    "register": "water_body_register.csv",
    "crosswalk": "monitoring_point_crosswalk.csv",
    "abstractions": "abstraction_point_register.csv",
    "manifest": "export_manifest.csv",
    "flow": "daily_mean_flow_2025_2026.csv",
    "sales": "pesticide_product_sales_returns_2025.csv",
    "programme": "monitoring_programme_2025_26.xlsx",
    "batches": "laboratory_batch_log.json",
}
EXPORT_COLS = ["lab_sample_id", "lab_code", "client_point_code", "sample_date", "sample_time",
               "date_received", "date_approved", "sample_type", "programme", "determinand_code",
               "method_code", "result", "units", "qualifier", "reporting_limit", "result_status",
               "version", "batch_id"]

ABBR = {"Cloonfad": "CF", "Glasheen": "GL", "Kilcrohane": "KC", "Owenbeg": "OB", "Ballynure": "BN",
        "Derryvore": "DV", "Tirneevin Stream": "TN", "Tirneevin": "TN", "Knockanure": "KN",
        "Carrowbeg": "CB", "Aghavannagh": "AG", "Moneycashen": "MC", "Drominagh": "DR",
        "Lisheenaguile": "LG"}


def client_code(lab, point, wb, section="catchment"):
    """the two laboratories number the authority's points their own way"""
    ab = ABBR.get(wb, "AR")
    tail = point.split("-")[1]
    n = int("".join(ch for ch in tail if ch.isdigit()) or "0")
    if tail[0].isalpha():
        ab = "AB"
    if lab == "TAS":
        base = "ARWA-%s-%02d" % (ab, n % 100)
    else:
        base = "AR%05d" % n
    if section == "services":
        base = base + ("-WS" if lab == "TAS" else "W")
    return base


def _fmt(v, det):
    if det in ("PH",):
        return "%.2f" % v
    if det in ("COND", "COLOUR"):
        return "%.0f" % v
    if det in ("TEMP", "DO", "SS", "ALK", "TURB"):
        return "%.1f" % v
    return "%.4f" % v


def build_rows():
    """every row of the export, in the order the LIMS writes it"""
    rng = random.Random(W.SEED + 97)
    rows = []
    sid = {"TAS": 183400, "KGL": 60710}
    batch_n = {"TAS": 0, "KGL": 0}

    def add(lab, point, place, wb, d, time, det, value, rl, detected, method, units, programme,
            sample_type="FIELD", status="A", version=1, section="catchment", approved=None):
        sid[lab] += 1
        yy = d.strftime("%y")
        bid = "%s-%s-%03d" % (lab, d.strftime("%Y%m"), 1 + (sid[lab] % 9))
        recv = d + dt.timedelta(days=1 + (sid[lab] % 2))
        appr = approved or (recv + dt.timedelta(days=9 + (sid[lab] % 11)))
        q = "<" if not detected else ""
        res = None if not detected else value
        rows.append({
            "lab_sample_id": "%s/%s/%07d" % (lab, yy, sid[lab]),
            "lab_code": lab, "client_point_code": client_code(lab, point, wb, section),
            "sample_date": d.isoformat(), "sample_time": time,
            "date_received": recv.isoformat(), "date_approved": appr.isoformat(),
            "sample_type": sample_type, "programme": programme, "determinand_code": det,
            "method_code": method, "result": ("" if res is None else _fmt(res, det)),
            "units": units, "qualifier": q,
            "reporting_limit": _fmt(rl, det) if rl else "", "result_status": status,
            "version": str(version), "batch_id": bid,
            "_point": point, "_wb": wb, "_value": value, "_det": detected, "_rl": rl,
            "_month": (d.year, d.month),
        })

    # --- the pesticide suite at the compliance points -------------------------------------------
    for wb in W.WATER_BODIES:
        name = wb[0]
        mrows = sim.mcpa_rows(wb)
        prng = random.Random(W.SEED + 400 + sum(ord(c) for c in name))
        for r in mrows:
            lab = r["lab"]
            d = sim.sample_day(prng, r["year"], r["month"], r["point"])
            tm = "%02d:%02d" % (9 + (sum(ord(c) for c in r["point"]) % 5), 5 * (prng.randrange(0, 11)))
            method = "TAS-LC-PHX-04" if lab == "TAS" else "KGL-GC-PHX-07"
            units = W.METHODS[method]["units"]
            scale = 1000.0 if units == "ng/l" else 1.0
            add(lab, r["point"], r["place"], name, d, tm, "MCPA", r["true"] * scale,
                r["rl"] * scale, r["det"], method, units, "PEST")
            for other in ("2,4-D", "MECOPROP", "TRICLOPYR"):
                f = {"2,4-D": 0.42, "MECOPROP": 0.31, "TRICLOPYR": 0.18}[other]
                v = r["true"] * f * (0.6 + prng.random() * 0.9)
                add(lab, r["point"], r["place"], name, d, tm, other, v * scale,
                    r["rl"] * scale, (v >= r["rl"]), method, units, "PEST")

    # --- the nutrient and physico-chemical suites ------------------------------------------------
    nrng = random.Random(W.SEED + 501)
    for r in R.nutrient_rows():
        d = sim.sample_day(nrng, r["year"], r["month"], r["point"])
        add(r["lab"], r["point"], r["place"], r["wb"], d, "10:%02d" % (5 * nrng.randrange(0, 11)),
            r["det"], r["value"], r["rl"], r["det_flag"], r["method"], r["units"], "NUT")
    prng2 = random.Random(W.SEED + 613)
    for r in R.physico_rows():
        d = sim.sample_day(prng2, r["year"], r["month"], r["point"])
        if r["half"]:
            d = d + dt.timedelta(days=14)
            if d.month != r["month"]:
                d = dt.date(r["year"], r["month"], 28)
        add(r["lab"], r["point"], r["place"], r["wb"], d, "11:%02d" % (5 * prng2.randrange(0, 11)),
            r["det"], r["value"], 0.0, True, r["method"], r["units"], "PHY")

    # --- the abstraction programme ---------------------------------------------------------------
    arng = random.Random(W.SEED + 727)
    for r in R.abstraction_rows():
        d = sim.sample_day(arng, r["year"], r["month"], r["point"])
        if r["half"]:
            d = d + dt.timedelta(days=14)
            if d.month != r["month"]:
                d = dt.date(r["year"], r["month"], 28)
        add(r["lab"], r["point"], r["place"], r["wb"], d,
            "08:%02d" % (5 * ((sum(ord(c) for c in r["point"]) + r["month"]) % 11)),
            r["det"], r["value"], r["rl"], True, r["method"], r["units"], "ABS",
            section="services" if r["dual"] else "catchment")

    # --- quality control rows --------------------------------------------------------------------
    field = [r for r in rows if r["sample_type"] == "FIELD"]
    qrng = random.Random(W.SEED + 839)
    for i, r in enumerate(field):
        rate = 9 if r["programme"] == "PEST" else 23
        if i % rate != 0:
            continue
        qc = W.QC_TYPES[(i // 23) % len(W.QC_TYPES)]
        v = 0.0 if qc.endswith("BLANK") else (
            r["_value"] + (r["_rl"] * 20 if r["_rl"] else 10.0) if qc == "MATRIX_SPIKE"
            else r["_value"] * (1.0 + qrng.gauss(0.0, 0.05)))
        detd = False if qc.endswith("BLANK") and r["_rl"] else True
        add(r["lab_code"], r["_point"], "", r["_wb"], dt.date.fromisoformat(r["sample_date"]),
            r["sample_time"], r["determinand_code"], v, r["_rl"], detd, r["method_code"],
            r["units"], r["programme"], sample_type=qc)

    # --- versions: the ortho-phosphate batches the laboratories reissued after the quarter --------
    q3 = [r for r in rows if r["determinand_code"] == "OPO4" and r["sample_type"] == "FIELD"
          and r["sample_date"][:7] in ("2025-10", "2025-11", "2025-12")]
    vrng = random.Random(W.SEED + 941)
    for i, r in enumerate(sorted(q3, key=lambda z: z["lab_sample_id"])):
        if i % 11 == 3:
            prov = dict(r)
            prov["result_status"], prov["version"] = "P", "1"
            prov["result"] = _fmt(r["_value"] * (1.0 + vrng.gauss(0.0, 0.08)), "OPO4")
            prov["date_approved"] = ""
            rows.append(prov)
            r["version"] = "2"
        elif i % 17 == 5:
            r["date_approved"] = (dt.date.fromisoformat(r["date_received"])
                                  + dt.timedelta(days=70)).isoformat()   # approved after the cut-off
        elif i % 19 == 7:
            r["result_status"] = "V"                                      # withdrawn on review
    return rows


def write_export(target, rows):
    path = os.path.join(target, FILES["export"])
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=EXPORT_COLS, lineterminator="\n")
        w.writeheader()
        for r in sorted(rows, key=lambda z: (z["lab_code"], z["lab_sample_id"], z["determinand_code"],
                                             z["version"])):
            w.writerow({k: r[k] for k in EXPORT_COLS})
    return path


def _csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        for r in rows:
            w.writerow(r)
    return path


def write_register(target):
    rows = []
    for (name, code, old, catch, area, target_mean, points) in W.WATER_BODIES:
        rows.append([code, name, old, catch, "%.1f" % area, "river", "2024-07-01", "compliance",
                     "in the authority's surface water compliance programme"])
    # three water bodies in the register that the compliance programme does not cover
    extra = [("IE_AR_18_1580", "Shanvally", "AR_0188", "Shanvally", 12.6),
             ("IE_AR_18_1690", "Bunlahy", "AR_0193", "Bunlahy", 9.4),
             ("IE_AR_18_1710", "Coolagh Lough", "AR_0201", "Coolagh", 2.8)]
    for (code, name, old, catch, area) in extra:
        rows.append([code, name, old, catch, "%.1f" % area,
                     "lake" if "Lough" in name else "river", "2024-07-01", "surveillance",
                     "surveillance network only, no monitoring programme at this water body"])
    return _csv(os.path.join(target, FILES["register"]),
                ["water_body_code", "water_body_name", "code_before_2024_revision", "catchment",
                 "area_km2", "category", "revision_effective", "network", "network_note"], rows)


def write_crosswalk(target, rows):
    seen = {}
    for r in rows:
        if not r["client_point_code"]:
            continue
        seen.setdefault((r["lab_code"], r["client_point_code"]), r)
    out = []
    pts = {}
    for (code, place, wb, wbcode, lab, pf) in R.compliance_points():
        pts[code] = (place, wb, wbcode)
    for (code, name, catch, lvl, lab) in W.ABSTRACTIONS:
        pts[code] = (name, catch, "")
    for (code, place, catch) in R.op_points():
        pts[code] = (place, catch, "")
    for (lab, ccode), r in sorted(seen.items()):
        p = r["_point"]
        place, wb, wbcode = pts.get(p, ("", r["_wb"], ""))
        section = "water services" if ccode.endswith(("-WS", "W")) and p.startswith("ARW-A") \
            else "catchment monitoring"
        out.append([p, W.RENUMBERED_FROM_2024.get(p, ""), place, wbcode, lab, ccode, section])
    return _csv(os.path.join(target, FILES["crosswalk"]),
                ["sample_point_code", "code_before_2024_revision", "sample_point_name",
                 "water_body_code", "lab_code", "client_point_code", "submitting_section"], out)


def write_abstractions(target):
    rows = [[code, name, catch, lab, "borehole" if "Borehole" in name or "Well" in name else
             ("spring" if "Spring" in name else "surface intake"), "fortnightly", "nitrate"]
            for (code, name, catch, lvl, lab) in W.ABSTRACTIONS]
    return _csv(os.path.join(target, FILES["abstractions"]),
                ["sample_point_code", "sample_point_name", "catchment", "lab_code", "source_type",
                 "nitrate_frequency", "programme"], rows)


def write_manifest(target, rows):
    dets = sorted(set(r["determinand_code"] for r in rows))
    dates = sorted(r["sample_date"] for r in rows)
    out = [["lims_results_export_2025_2026.csv", str(len(rows)), dates[0], dates[-1],
            str(len(dets)), ";".join(dets), W.AS_OF.isoformat(),
            "all result versions as held on the extract date, both laboratories"]]
    return _csv(os.path.join(target, FILES["manifest"]),
                ["file_name", "row_count", "earliest_sample_date", "latest_sample_date",
                 "determinand_count", "determinands", "extracted", "scope"], out)


def write_flow(target):
    """hydrometric daily mean flow: context the authority holds, no part of the assessment"""
    rng = random.Random(W.SEED + 1301)
    gauges = [("AR18001", "Cloonfad Bridge"), ("AR18004", "Owenbeg at Drumlosh"),
              ("AR18007", "Ballynure at Tooreen"), ("AR18011", "Glasheen at Rathvilly"),
              ("AR18014", "Carrowbeg at Ballinaboy")]
    rows = []
    d = W.YEAR_FROM
    while d <= W.YEAR_TO:
        for (g, name) in gauges:
            base = 1.4 + (sum(ord(c) for c in g) % 23) / 6.0
            wet = 1.0 + (0.85 if d.month in (11, 12, 1, 2) else -0.45 if d.month in (6, 7, 8) else 0.1)
            rows.append([g, name, d.isoformat(), "%.3f" % max(0.02, base * wet * (1.0 + rng.gauss(0, 0.33))),
                         "good" if rng.random() > 0.03 else "estimated"])
        d += dt.timedelta(days=1)
    return _csv(os.path.join(target, FILES["flow"]),
                ["station_code", "station_name", "date", "daily_mean_flow_m3s", "quality"], rows)


def write_sales(target):
    """merchant returns of phenoxy herbicide product, lodged under the plant protection regulations"""
    rng = random.Random(W.SEED + 1409)
    merchants = ["Araglen Farm Supplies", "Kilgarron Agri Store", "Boherboy Co-operative",
                 "Tooreen Agricultural Merchants", "Lisnagorman Farm Services",
                 "Shanacloon Growers Supply"]
    rows = []
    for mn in merchants:
        for (y, m) in W.MONTHS:
            for prod, ai in (("Kilrush 750 SL", "MCPA"), ("Rushmaster XL", "MCPA"),
                             ("Bracklyn D", "2,4-D"), ("Thistletop", "MECOPROP")):
                season = 1.0 + (2.6 if m in (5, 6) else 1.1 if m in (4, 7) else -0.8)
                litres = max(0.0, 180.0 * season * (0.5 + (sum(ord(c) for c in mn) % 17) / 12.0)
                             * (1.0 + rng.gauss(0, 0.4)))
                rows.append([mn, "%04d-%02d" % (y, m), prod, ai, "%.0f" % litres,
                             "pasture, rushes and docks"])
    return _csv(os.path.join(target, FILES["sales"]),
                ["merchant", "month", "product_name", "active_substance", "litres_sold",
                 "label_use"], rows)


def write_batches(target, rows):
    """the laboratories' batch log, as the contract requires it"""
    out = {"district": W.DISTRICT, "compiled": W.AS_OF.isoformat(), "batches": []}
    seen = {}
    for r in rows:
        b = seen.setdefault(r["batch_id"], {"batch_id": r["batch_id"], "lab_code": r["lab_code"],
                                            "results": 0, "determinands": set(),
                                            "earliest_sample": r["sample_date"],
                                            "latest_sample": r["sample_date"]})
        b["results"] += 1
        b["determinands"].add(r["determinand_code"])
        b["earliest_sample"] = min(b["earliest_sample"], r["sample_date"])
        b["latest_sample"] = max(b["latest_sample"], r["sample_date"])
    for k in sorted(seen):
        b = seen[k]
        b["determinands"] = sorted(b["determinands"])
        out["batches"].append(b)
    path = os.path.join(target, FILES["batches"])
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(out, fh, indent=1, sort_keys=False)
        fh.write("\n")
    return path


def write_programme(target, rows):
    """the monitoring programme workbook, in the authority's own wide layout"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Programme 2025-26"
    ws["A1"] = "%s: surface water monitoring programme" % W.AUTHORITY
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Reporting year 1 April 2025 to 31 March 2026. Required rounds per point and suite."
    ws["A4"] = "Sample point"
    ws["B4"] = "Name"
    ws["C4"] = "Water body"
    ws["D4"] = "Laboratory"
    ws["E4"] = "Required rounds"
    ws["E5"] = "Pesticide suite"
    ws["F5"] = "Nutrient suite"
    ws["G5"] = "Physico-chemical"
    ws["H4"] = "In programme from"
    ws["I4"] = "In programme to"
    ws["J4"] = "Notes"
    ws.merge_cells("E4:G4")
    for c in ("A4", "B4", "C4", "D4", "E4", "H4", "I4", "J4", "E5", "F5", "G5"):
        ws[c].font = Font(bold=True)
        ws[c].alignment = Alignment(horizontal="center", wrap_text=True)
    row = 6
    for (code, place, wbname, wbcode, lab, pf) in R.compliance_points():
        cal = len(sim.point_calendar(code))
        frm = W.POINT_FROM.get(code)
        to = W.POINT_TO.get(code)
        note = ""
        if frm:
            note = "added at the catchment partnership's request"
        elif to:
            note = "access withdrawn by the landowner"
        ws.cell(row=row, column=1, value=code)
        ws.cell(row=row, column=2, value=place)
        ws.cell(row=row, column=3, value=wbname)
        ws.cell(row=row, column=4, value=W.LABS[lab]["name"])
        ws.cell(row=row, column=5, value=12 if not (frm or to) else 6)
        ws.cell(row=row, column=6, value=12)
        ws.cell(row=row, column=7, value=24)
        ws.cell(row=row, column=8, value=(frm or W.YEAR_FROM).isoformat())
        ws.cell(row=row, column=9, value=(to or W.YEAR_TO).isoformat())
        ws.cell(row=row, column=10, value=note or None)
        row += 1
    ws2 = wb.create_sheet("Rounds not sampled")
    ws2["A1"] = "Pesticide rounds the programme required and the field team did not complete"
    ws2["A1"].font = Font(bold=True)
    ws2["A3"], ws2["B3"], ws2["C3"], ws2["D3"] = "Sample point", "Year", "Month", "Reason"
    for c in ("A3", "B3", "C3", "D3"):
        ws2[c].font = Font(bold=True)
    reasons = ["river in spate, no safe access", "gate locked, landowner away",
               "ice on the bank, round abandoned", "road closed for resurfacing",
               "sampler diverted to a pollution incident"]
    for i, (code, y, m) in enumerate(W.MISSED):
        ws2.cell(row=4 + i, column=1, value=code)
        ws2.cell(row=4 + i, column=2, value=y)
        ws2.cell(row=4 + i, column=3, value=m)
        ws2.cell(row=4 + i, column=4, value=reasons[i % len(reasons)])
    ws3 = wb.create_sheet("Operational network")
    ws3["A1"] = "Operational and surveillance points: nutrient and physico-chemical suites only"
    ws3["A1"].font = Font(bold=True)
    ws3["A3"], ws3["B3"], ws3["C3"] = "Sample point", "Name", "Catchment"
    for c in ("A3", "B3", "C3"):
        ws3[c].font = Font(bold=True)
    for i, (code, place, catch) in enumerate(R.op_points()):
        ws3.cell(row=4 + i, column=1, value=code)
        ws3.cell(row=4 + i, column=2, value=place)
        ws3.cell(row=4 + i, column=3, value=catch)
    for w, widths in ((ws, [14, 34, 18, 30, 11, 11, 11, 18, 18, 42]),
                      (ws2, [14, 8, 8, 40]), (ws3, [14, 34, 18])):
        for i, wd in enumerate(widths, start=1):
            w.column_dimensions[chr(64 + i)].width = wd
    path = os.path.join(target, FILES["programme"])
    wb.save(path)
    return path
