"""The golden deliverable: leaver_destinations_2026_quality_flags.xlsx."""
import os

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import params as P
import postprocess

XLSX = "leaver_destinations_2026_quality_flags.xlsx"
HEAD = PatternFill("solid", fgColor="DCE6F0")
FILL = {"Suppressed": PatternFill("solid", fgColor="F4CCCC"), "Caution": PatternFill("solid", fgColor="FFF2CC"),
        "Standard": PatternFill("solid", fgColor="FFFFFF")}
THIN = Border(bottom=Side(style="thin", color="BBBBBB"))
STATNAME = {"neet": "NEET rate", "he": "Higher-education entry rate"}


def _head(ws, row, labels, widths):
    for j, (lab, w) in enumerate(zip(labels, widths), 1):
        c = ws.cell(row=row, column=j, value=lab)
        c.font = Font(bold=True); c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = w


def write(folder, A):
    wb = Workbook()
    ws = wb.active; ws.title = "Flags"
    ws["A1"] = ("Leaver destinations 2026: %d district rates suppressed, %d published with the standard flag, %d under caution"
                % (len(A["lists"]["Suppressed"]), len(A["lists"]["Standard"]), len(A["lists"]["Caution"])))
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("QS-4 v2.0. Flag on the size of the sensitivity shift: under 1.0 point standard, 1.0 to under 3.0 caution, "
                "3.0 or more suppressed. Frame extract of 10 November 2026; 2026 district map.")
    ws["A2"].font = Font(italic=True, size=9)
    _head(ws, 4, ["District code", "District", "Statistic", "Headline rate (%)", "Sensitivity shift (points)", "Flag"],
          [14, 18, 28, 18, 24, 12])
    r = 5
    for c in A["flags"]:
        ws.cell(row=r, column=1, value=c["district"]); ws.cell(row=r, column=2, value=c["name"])
        ws.cell(row=r, column=3, value=STATNAME[c["stat"]])
        x = ws.cell(row=r, column=4, value=c["rate"]); x.number_format = "0.0"
        y = ws.cell(row=r, column=5, value=c["shift"]); y.number_format = "+0.00;-0.00"
        z = ws.cell(row=r, column=6, value=c["flag"]); z.fill = FILL[c["flag"]]; z.font = Font(bold=c["flag"] != "Standard")
        for j in range(1, 7):
            ws.cell(row=r, column=j).border = THIN
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Suppressed").font = Font(bold=True)
    ws.cell(row=r, column=2, value="; ".join(A["lists"]["Suppressed"]))
    ws.cell(row=r + 1, column=1, value="Standard").font = Font(bold=True)
    ws.cell(row=r + 1, column=2, value="; ".join(A["lists"]["Standard"]))
    ws.cell(row=r + 2, column=1, value="Caution").font = Font(bold=True)
    ws.cell(row=r + 2, column=2, value="the other %d statistics" % len(A["lists"]["Caution"]))
    ws.cell(row=r + 4, column=1, value=("A suppressed statistic's headline rate is withheld from the published table (QS-4 section 4); "
                                       "it is shown here for the record.")).font = Font(italic=True, size=9)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "A4:F%d" % (4 + len(A["flags"]))

    wb2 = wb.create_sheet("Attainment bands")
    wb2["A1"] = "Leaver destinations 2026: attainment bands across the island (bands on final examination points)"
    wb2["A1"].font = Font(bold=True, size=12)
    _head(wb2, 3, ["Attainment band", "Leavers", "NEET rate (%)", "Higher-education entry rate (%)"], [16, 12, 16, 30])
    for i, b in enumerate(A["bands"]):
        wb2.cell(row=4 + i, column=1, value=b["band"])
        wb2.cell(row=4 + i, column=2, value=b["leavers"]).number_format = "#,##0"
        wb2.cell(row=4 + i, column=3, value=b["neet"]).number_format = "0.0"
        wb2.cell(row=4 + i, column=4, value=b["he"]).number_format = "0.0"
    wb2.cell(row=8, column=1, value="All bands").font = Font(bold=True)
    wb2.cell(row=8, column=2, value=sum(b["leavers"] for b in A["bands"])).number_format = "#,##0"

    ws3 = wb.create_sheet("Follow-up departure")
    ws3["A1"] = "2025 nonresponse follow-up: found destinations against allocations, by followed-up district and band"
    ws3["A1"].font = Font(bold=True, size=12)
    ws3["A2"] = ("Cases resolved and still nonrespondents in the 2025 final file, one per case: %s. On the logit of the allocation "
                 "the found destinations have slope %.2f (NEET) and %.2f (higher education), one odds ratio throughout; the cells "
                 "scatter around it without trend, while the points gap rises with the allocation." % (format(A["n_cases"], ","),
                 A["slope"]["neet"], A["slope"]["he"]))
    ws3["A2"].font = Font(italic=True, size=9)
    _head(ws3, 4, ["Statistic", "District", "Band", "Cases", "Mean allocation (%)", "Found (%)", "Gap (points)",
                   "Ratio of rates", "Odds ratio"], [26, 16, 8, 8, 18, 11, 13, 14, 11])
    r = 5
    for s in ("neet", "he"):
        for c in A["cells"][s]:
            vals = [STATNAME[s], c["name"], c["band"], c["n"], round(c["alloc"] * 100, 2), round(c["found"] * 100, 2),
                    round(c["gap"] * 100, 2), round(c["ratio"], 3), round(c["odds_ratio"], 3)]
            fmts = [None, None, None, "0", "0.0", "0.0", "+0.0;-0.0", "0.00", "0.00"]
            for j, (v, f) in enumerate(zip(vals, fmts), 1):
                x = ws3.cell(row=r, column=j, value=v)
                if f:
                    x.number_format = f
            r += 1
        ws3.cell(row=r, column=1, value=STATNAME[s] + ": pooled odds ratio").font = Font(bold=True)
        ws3.cell(row=r, column=9, value=A["odds_ratio"][s]).number_format = "0.00"
        r += 2
    ws3.cell(row=r, column=1, value="How each form of departure fits the follow-up's cases (log-likelihood; higher fits better)").font = Font(bold=True)
    r += 1
    _head(ws3, r, ["Form of departure", "NEET", "Higher education"], [26, 16, 8])
    r += 1
    for k in A["ll"]["neet"]:
        ws3.cell(row=r, column=1, value=k)
        ws3.cell(row=r, column=2, value=round(A["ll"]["neet"][k], 1)).number_format = "0.0"
        ws3.cell(row=r, column=3, value=round(A["ll"]["he"][k], 1)).number_format = "0.0"
        r += 1
    ws3.freeze_panes = "A5"

    ws4 = wb.create_sheet("Notes")
    notes = [
        "Basis",
        "Released rate: returned forms in the destination plus nonrespondents' allocations, over every leaver on the 2026 frame "
        "in the district (QS-4 section 2).",
        "Sensitivity shift: each nonrespondent's allocation moved as the 2025 follow-up found followed-up nonrespondents "
        "departing from theirs, worked nonrespondent by nonrespondent and summed to the district, minus the released rate "
        "(QS-4 section 3).",
        "The follow-up's departure is one odds ratio across its five districts and four bands (sheet Follow-up departure: "
        "slope %.2f and %.2f on the logit of the allocation), so each nonrespondent's allocation is moved on that scale: odds "
        "multiplied by %.2f for NEET and %.2f for higher education." % (A["slope"]["neet"], A["slope"]["he"],
                                                                      A["odds_ratio"]["neet"], A["odds_ratio"]["he"]),
        "Flags are set on the unrounded shift (QS-4 section 5); rates are shown to one decimal place.",
        "Returned forms",
        "Online returns made on build 7 of the form (before the 2026 edition went live) carry the 2025 edition's activity "
        "codes and are read with that edition's order, as the 2025 final file's processed destinations show it.",
        "Attainment bands are set on the final examination points after appeals (methodology section 5), not on the band "
        "supplied with the September leaver return.",
    ]
    for i, t in enumerate(notes, 1):
        c = ws4.cell(row=i, column=1, value=t)
        if t in ("Basis", "Returned forms"):
            c.font = Font(bold=True)
        else:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws4.column_dimensions["A"].width = 120
    path = os.path.join(folder, XLSX)
    wb.save(path)
    postprocess.scrub(path, P.PEOPLE["head"], "2026-11-13")
    postprocess.finish_office(path, "2026-11-12T09:10:00Z", "2026-11-13T15:45:00Z", P.ORG)
    return path
