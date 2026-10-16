"""task233 documents: the quality standard QS-4, the survey methodology and the 2025 follow-up's field report (pdf),
the school register (xlsx), the data dictionary and the readme (md)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess

OFFICE = P.ORG
DEPT = "Department of Education"


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=13, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=14, leading=18, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=14, spaceBefore=8,
                        spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10, textColor=colors.HexColor("#444444"))
    return dict(body=body, h1=h1, h2=h2, small=small)


def _footer(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(18 * mm, 10 * mm, text)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _pdf(path, story, title, author, footer, stamp):
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title=title, author=author, creator=author, invariant=1)
    doc.build(story, onFirstPage=_footer(footer), onLaterPages=_footer(footer))
    postprocess.scrub(path, author, stamp)


def _table(rows, widths, size=8.0):
    t = Table(rows, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", size), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", size),
                           ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6ECF2")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


# ------------------------------------------------------------------ QS-4
def standard(path):
    S = _styles(); b, h1, h2, sm = S["body"], S["h1"], S["h2"], S["small"]
    st = [Paragraph("QS-4 Quality flags for district destination statistics", h1),
          Paragraph("%s, %s. Version 2.0, approved 14 September 2026 by %s, survey director, and agreed with the %s. "
                    "Replaces version 1.3 (2023), which flagged a district rate on its nonresponse share alone." % (OFFICE, P.DIVISION, P.PEOPLE["head"], DEPT), sm),
          Spacer(1, 4),
          Paragraph("1. Scope", h2),
          Paragraph("This standard sets the quality flag carried by the NEET rate and the higher-education entry rate published "
                    "for each education district in the %s's school-leaver destinations release, from the 2026 release. Districts are "
                    "those of the district map in force on the census date." % DEPT, b),
          Paragraph("2. Definitions", h2),
          Paragraph("A <i>leaver</i> is a pupil on the release frame: a leaver of the school year from a state-funded post-primary "
                    "school. A <i>nonrespondent</i> is a leaver from whom no form was returned by the close of fieldwork. Each "
                    "nonrespondent carries an <i>allocation</i>: the share of each destination the destination model gives that "
                    "leaver (methodology, section 4). The <i>headline rate</i> of a destination in a district is the number of "
                    "leavers whose returned form places them in it plus the nonrespondents' allocations to it, divided by all "
                    "leavers on the frame in the district, as a percentage. It is computed for every district whatever its flag.", b),
          Paragraph("3. Sensitivity rate and sensitivity shift", h2),
          Paragraph("Allocation assumes that nonrespondents go where respondents like them go. The sensitivity rate tests that "
                    "assumption against the evidence: it is the headline rate recomputed with each nonrespondent's "
                    "allocation moved as the most recent nonresponse follow-up found followed-up nonrespondents' "
                    "destinations departing from their allocations. It is worked for each nonrespondent in turn and summed to the "
                    "district. The sensitivity shift is the sensitivity rate minus the headline rate, in percentage points.", b),
          Paragraph("4. Flags", h2),
          _table([["Size of the sensitivity shift", "Flag", "In the release"],
                  ["under 1.0 percentage point", "Standard", "published without comment"],
                  ["1.0 to under 3.0 points", "Caution", "published with the caution note"],
                  ["3.0 points or more", "Suppressed", "headline rate withheld; the table shows the suppression mark"]],
                 [62 * mm, 30 * mm, 80 * mm]),
          Spacer(1, 6),
          Paragraph("5. Rounding", h2),
          Paragraph("The flag is set on the unrounded shift. Headline rates are shown to one decimal place.", b),
          Paragraph("6. The follow-up", h2),
          Paragraph("The departure is measured on the follow-up's resolved cases whose leaver remained a nonrespondent in that "
                    "year's final file, each case counted once, against the allocation that leaver carries in the final file. The "
                    "follow-up is run in a sample of districts and stands for all of them: every district's sensitivity rate is "
                    "worked from it.", b),
          Paragraph("7. Sign-off", h2),
          Paragraph("The survey director signs the flags after the allocation run and sends them to the %s with the release "
                    "tables. A flag once published is not changed before the following year's release." % DEPT, b)]
    _pdf(path, st, "QS-4 Quality flags for district destination statistics", P.PEOPLE["head"],
         "%s | QS-4 v2.0 | %s" % (OFFICE, P.DIVISION), "2026-09-14")


# ------------------------------------------------------------------ methodology
def methodology(path, counts):
    S = _styles(); b, h1, h2, sm = S["body"], S["h1"], S["h2"], S["small"]
    st = [Paragraph("Leaver Destinations Survey 2026: methodology", h1),
          Paragraph("%s, %s. Prepared by %s, survey methodologist, 10 November 2026." % (OFFICE, P.DIVISION, P.PEOPLE["method"]), sm),
          Spacer(1, 4),
          Paragraph("1. The survey and its frame", h2),
          Paragraph("Every leaver of the 2025-26 school year from a state-funded post-primary school is asked where they were on "
                    "the census date, 1 October 2026. The frame is built from the schools' September leaver returns: %s leavers "
                    "in %d schools across the twelve education districts. Each leaver is asked once; there is no sampling and no "
                    "weighting." % (format(counts["frame26"], ","), counts["schools"], ), b),
          Paragraph("2. Destinations", h2),
          Paragraph("The release reports five destinations: higher education (HE, a course on or off the island), further "
                    "education (FE), apprenticeship or other training (APP), employment (EMP), and not in education, employment or "
                    "training (NEET: unemployed and seeking work, or not working or studying). Each returned form records one "
                    "activity, which the activity code list maps to its destination.", b),
          Paragraph("3. Fieldwork", h2),
          Paragraph("Paper forms (LDS-1) went out with a reply envelope on 21 September 2026; the online form (LDS-2) opened at "
                    "09:00 on Thursday 1 October. Reminders went out on 8 and 19 October. Fieldwork closed at 17:00 on Friday "
                    "30 October 2026. Paper forms are keyed by the Unit's data capture team; online returns are written to the "
                    "frame as submitted. Times on the frame are in UTC.", b),
          Paragraph("4. Allocation of nonrespondents", h2),
          Paragraph("A form not returned by the close leaves the leaver's destination unknown. The release allocates each "
                    "nonrespondent across the five destinations in the proportions found among the previous year's respondents "
                    "in the same district and attainment band (the previous year's final file, restated on the current district "
                    "map), because a year's forms are fully processed only when its final file closes at the end of March. The "
                    "allocation run of 9 November 2026 wrote each nonrespondent's five shares to the frame.", b),
          Paragraph("5. Attainment bands", h2),
          Paragraph("Leavers are banded on their state examination points as finalised after the appeals process (revised results "
                    "issued 20 October 2026): AB1 480 points or more, AB2 360 to 479, AB3 240 to 359, AB4 under 240.", b),
          Paragraph("6. District map", h2),
          Paragraph("A boundary order effective 1 September 2026 moved two schools to a neighbouring district. The school register "
                    "gives each school's district on the 2025 and the 2026 map; the 2026 release uses the 2026 map.", b),
          Paragraph("7. The release", h2),
          Paragraph("The %s's release of 10 December 2026 publishes, for each district and for each attainment band across the "
                    "island, the number of leavers and the rate of each destination, with the quality flags of QS-4 on each "
                    "district's NEET rate and higher-education entry rate." % DEPT, b),
          Paragraph("8. Files in this pack", h2),
          Paragraph("The 2026 frame as extracted on 10 November 2026, after the allocation run; the 2025 final file of 31 March "
                    "2026, which carries the 2025 forms as processed (with each respondent's destination) and the 2025 "
                    "allocations of the leavers who never returned a form; the 2025 nonresponse follow-up's outcomes, its case "
                    "crosswalk and its field report.", b)]
    _pdf(path, st, "Leaver Destinations Survey 2026: methodology", P.PEOPLE["method"],
         "%s | Leaver Destinations Survey 2026 | methodology" % OFFICE, "2026-11-10")


# ------------------------------------------------------------------ follow-up report
def followup_report(path, comp):
    S = _styles(); b, h1, h2, sm = S["body"], S["h1"], S["h2"], S["small"]
    rows = [["District", "Code", "Cases drawn", "Resolved", "Not resolved"]]
    for r in comp:
        rows.append([r["name"], r["code"], format(r["drawn"], ","), format(r["resolved"], ","), str(r["unresolved"])])
    rows.append(["All five", "", format(sum(r["drawn"] for r in comp), ","), format(sum(r["resolved"] for r in comp), ","),
                 str(sum(r["unresolved"] for r in comp))])
    codes = [["Code", "Outcome found", "Release destination"],
             ["F01", "Higher education course (on or off the island)", "HE"], ["F02", "Further education course", "FE"],
             ["F03", "Apprenticeship", "APP"], ["F04", "Other training programme", "APP"], ["F05", "In employment", "EMP"],
             ["F06", "Unemployed and seeking work", "NEET"], ["F07", "Not working or studying (caring, illness, other)", "NEET"],
             ["F99", "Not resolved", "none"]]
    st = [Paragraph("Nonresponse follow-up 2025: field report", h1),
          Paragraph("%s, Field Operations. %s, field manager. 16 March 2026." % (OFFICE, P.PEOPLE["field"]), sm),
          Spacer(1, 4),
          Paragraph("1. Purpose", h2),
          Paragraph("The follow-up establishes where leavers who did not return the 2025 destinations form actually went, so "
                    "that QS-4's sensitivity measure rests on observed destinations rather than on assumption.", b),
          Paragraph("2. Design", h2),
          Paragraph("The follow-up ran in five districts chosen across the island's range of area deprivation: Upper Valley, Northern Plains, Riverside, City East and Mill Towns. On %s the field team drew a simple random sample of 88 in every "
                    "100 of the leavers in those districts from whom no form had been returned when fieldwork closed (17:00 on "
                    "24 October 2025). Fieldwork ran from 3 November 2025 to 27 February 2026." % "31 October 2025", b),
          Paragraph("3. Routes and recording", h2),
          Paragraph("Field officers resolved a case through an adviser visit, a telephone interview, or a record match with the "
                    "benefits register, college enrolment records or employer records. Where two routes found the same case, "
                    "both are recorded in the outcomes file with the date of each; in every such case the routes agreed, and the "
                    "case has one outcome. Cases are keyed by the field system's case number; the crosswalk gives each case's "
                    "leaver reference as the field system holds it.", b),
          Paragraph("4. Completion", h2),
          _table(rows, [40 * mm, 18 * mm, 30 * mm, 30 * mm, 30 * mm]),
          Spacer(1, 4),
          Paragraph("5. Outcome codes", h2),
          _table(codes, [16 * mm, 110 * mm, 40 * mm])]
    _pdf(path, st, "Nonresponse follow-up 2025: field report", P.PEOPLE["field"],
         "%s | Nonresponse follow-up 2025 | field report" % OFFICE, "2026-03-16")


# ------------------------------------------------------------------ school register
def register(path, reg):
    wb = Workbook(); ws = wb.active; ws.title = "Schools"
    hdr = ["school_code", "school_name", "district_2025", "district_2026"]
    ws.append(hdr)
    for r in reg.itertuples():
        ws.append([r.school_code, r.school_name, r.district_2025, r.district_2026])
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="E6ECF2")
    for i, w in enumerate([13, 44, 14, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    wd = wb.create_sheet("Districts")
    wd.append(["district_code", "district"])
    for code, name in sorted(zip(P.CODE, P.NAME)):
        wd.append([code, name])
    for c in wd[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="E6ECF2")
    wd.column_dimensions["A"].width = 14; wd.column_dimensions["B"].width = 22
    wb.save(path)
    postprocess.scrub(path, P.PEOPLE["method"], "2026-09-02")
    postprocess.finish_office(path, "2026-08-31T10:12:00Z", "2026-09-02T15:40:00Z", OFFICE)


def release_tables(path, t25):
    wb = Workbook(); ws = wb.active; ws.title = "Districts 2025"
    ws.append(["Leaver destinations 2025: district rates as published on 11 December 2025 (2025 district map)"])
    ws["A1"].font = Font(bold=True)
    ws.append([])
    ws.append(["district_code", "district", "leavers", "neet_rate", "he_entry_rate"])
    for c in ws[3]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="E6ECF2")
    for r in t25.itertuples():
        ws.append([r.district_code, r.district, r.leavers, r.neet_rate, r.he_entry_rate])
    for row in ws.iter_rows(min_row=4, min_col=4, max_col=5):
        for c in row:
            c.number_format = "0.0"
    for i, w in enumerate([14, 20, 10, 11, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(path)
    postprocess.scrub(path, P.PEOPLE["head"], "2025-12-11")
    postprocess.finish_office(path, "2025-12-02T09:30:00Z", "2025-12-11T08:15:00Z", OFFICE)
