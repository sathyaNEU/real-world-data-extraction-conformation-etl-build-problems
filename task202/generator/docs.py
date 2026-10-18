"""task202 documents: the three monthly claim return workbooks (the band columns renamed and a header
block added from the 2026/27 year), the scheme's operating policy (pdf), the prior year's published
partner summary (pdf), the data dictionary (md), the export manifest (csv) and the request (eml)."""
import calendar
import contextlib
import datetime as dt
import os
import re
import shutil
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import records as RC
import sim


PRODUCER = "%s document services" % P.FOUNDATION_LEGAL
APPNAME = "Microsoft Excel"


def epoch(y, m, d, hh=9, mm=0):
    return calendar.timegm((y, m, d, hh, mm, 0, 0, 0, 0))


def org_canvas():
    """The organisation's own document stamp, so no shipped file carries the writing library's
    signature."""
    from reportlab.pdfgen import canvas as _c

    class OrgCanvas(_c.Canvas):
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self.setProducer(PRODUCER)

    return OrgCanvas


def finalise_pdf(path, producer=None):
    """Put the organisation's name where the writing library signed the file, in place and at the
    same byte length so every cross-reference offset stays valid. reportlab writes its own name into
    the header comment, the trailer digest comment and the producer slot whatever the document's
    properties say."""
    producer = producer or PRODUCER
    raw = open(path, "rb").read()
    before = _pdf_text(path)
    out = raw
    for tok in (b"ReportLab Generated PDF document (opensource)",
                b"% ReportLab generated PDF document -- digest (opensource)"):
        if tok in out:
            lead = b"% " if tok.startswith(b"% ") else b""
            name = (lead + ("%s filed document" % P.FOUNDATION_LEGAL).encode())[:len(tok)]
            out = out.replace(tok, name.ljust(len(tok), b" "))
    m = re.search(rb"/Producer \(([^)]*)\)", out)
    if m:
        slot = len(m.group(1))
        out = out[:m.start(1)] + producer.encode()[:slot].ljust(slot, b" ") + out[m.end(1):]
    assert len(out) == len(raw), (len(out), len(raw))
    open(path, "wb").write(out)
    assert _pdf_text(path) == before, "the substitution changed the document's own text"


def _pdf_text(path):
    from pypdf import PdfReader
    return "\n".join((pg.extract_text() or "") for pg in PdfReader(path).pages)


@contextlib.contextmanager
def stamped(ts):
    """The date the document states about itself, which is the date the organisation filed it and
    never the build clock."""
    old = os.environ.get("SOURCE_DATE_EPOCH")
    os.environ["SOURCE_DATE_EPOCH"] = str(int(ts))
    try:
        yield
    finally:
        if old is None:
            os.environ.pop("SOURCE_DATE_EPOCH", None)
        else:
            os.environ["SOURCE_DATE_EPOCH"] = old


def finalise_xlsx(path, stamp):
    """Repack the workbook with the export's own date on every entry and the spreadsheet
    application in docProps, entry names and order untouched. openpyxl stamps each zip entry with
    the build clock, which is both an anachronism and a byte difference between two builds."""
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        iso = ("%04d-%02d-%02dT%02d:%02d:%02dZ" % stamp).encode()
        for i, data in items:
            if i.filename == "docProps/app.xml":
                data = re.sub(rb"<Application>[^<]*</Application>",
                              b"<Application>" + APPNAME.encode() + b"</Application>", data)
            if i.filename == "docProps/core.xml":
                # openpyxl writes the build clock into dcterms:modified whatever the workbook's own
                # properties say, which is both an anachronism and a byte difference between builds
                data = re.sub(rb"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", iso, data)
            zi = zipfile.ZipInfo(i.filename, date_time=stamp)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = i.external_attr
            out.writestr(zi, data)
    shutil.move(tmp, path)


def addr(k):
    n, u, _ = P.PEOPLE[k]
    return "%s <%s@%s>" % (n, u, P.DOMAIN)


def styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=9.4, leading=13.2,
                          spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=13, leading=16,
                        spaceAfter=3)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.2,
                        leading=13.5, spaceBefore=9, spaceAfter=3, keepWithNext=1)
    small = ParagraphStyle("s", parent=body, fontSize=7.9, leading=10.4,
                           textColor=colors.HexColor("#3f3f3f"))
    cell = ParagraphStyle("c", fontName="Helvetica", fontSize=8.3, leading=10.1)
    return body, h1, h2, small, cell


def grid(rows, widths, cell, align_right=()):
    t = Table([[Paragraph(str(c), cell) for c in r] for r in rows], colWidths=widths, repeatRows=1)
    st = [("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#97a0ab")),
          ("VALIGN", (0, 0), (-1, -1), "TOP"),
          ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e4e9f0"))]
    for c in align_right:
        st.append(("ALIGN", (c, 1), (c, -1), "RIGHT"))
    t.setStyle(TableStyle(st))
    return t


def footer(text):
    def f(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.2)
        canvas.setFillColor(colors.HexColor("#5a5a5a"))
        canvas.drawString(20 * mm, 11 * mm, text)
        canvas.drawRightString(A4[0] - 20 * mm, 11 * mm, "Page %d" % d.page)
        canvas.restoreState()
    return f


def doc(path, title, author, subject):
    return SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
                             topMargin=17 * mm, bottomMargin=19 * mm, title=title, author=author,
                             creator=author, subject=subject, invariant=1)


# ------------------------------------------------------------------ the claim return workbooks
def _fill(ws, r0, header, rows, widths):
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="E4E9F0")
    for j, h in enumerate(header, start=1):
        c = ws.cell(row=r0, column=j, value=h)
        c.font = bold
        c.fill = fill
        c.alignment = Alignment(horizontal="left")
    for i, row in enumerate(rows, start=r0 + 1):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=v)
            if isinstance(v, float):
                c.number_format = "#,##0.00"
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)


def claim_returns(path, R, months, new_layout, fy, recovery=None):
    wb = Workbook()
    ws = wb.active
    rows = R.returns_rows(months)
    if recovery is not None:
        mo, code, b, n, tot, _ = recovery
        rows = rows + [(mo, code, b, n, round(tot, 2), "FINAL")]
    if new_layout:
        ws.title = "ClaimReturns"
        ws["A1"] = "%s - %s: monthly claim returns" % (P.FOUNDATION, P.SCHEME)
        ws["A2"] = "Fiscal year %s (claims to the month shown)" % fy
        ws["A3"] = "Extract taken %s by the %s" % (P.EXTRACT_DATE.strftime("%d %B %Y"), P.OFFICE)
        ws["A4"] = "Amounts in pesos. One row per partner, household-size band and claim month."
        for r in range(1, 5):
            ws.cell(row=r, column=1).font = Font(bold=(r == 1), italic=(r > 2))
        head = ["claim_month", "partner_code", "household_size_band", "awards",
                "total_claimed_php", "return_status"]
        data = [(RC.mlabel(mo, True), code, RC.legacy_band(b), n, float(tot), st)
                for mo, code, b, n, tot, st in rows]
        data.sort(key=lambda r: (months.index(_unm(r[0])), r[1], r[2], r[5]))
        _fill(ws, 6, head, data, [13, 15, 20, 9, 18, 14])
    else:
        ws.title = "Returns"
        head = ["return_month", "partner_code", "hh_band", "award_count", "claim_total_php"]
        data = [(RC.mlabel(mo), code, RC.band_code(b), n, float(tot))
                for mo, code, b, n, tot, st in rows]
        data.sort(key=lambda r: (r[0], r[1], r[2]))
        _fill(ws, 1, head, data, [14, 15, 10, 13, 18])
    wb.properties.creator = P.OFFICE
    wb.properties.lastModifiedBy = P.OFFICE
    wb.properties.title = "%s claim returns %s" % (P.SCHEME_SHORT, fy)
    wb.properties.created = dt.datetime(2027, 4, 30, 9, 15)
    wb.properties.modified = dt.datetime(2027, 4, 30, 9, 15)
    wb.save(path)
    finalise_xlsx(path, (2027, 4, 30, 9, 15, 0))


_MON = {dt.date(2000, m, 1).strftime("%b"): m for m in range(1, 13)}


def _unm(lab):
    a, b = lab.split("-")
    return (int(b), _MON[a])


# ------------------------------------------------------------------ the operating policy
def policy(path, R):
    body, h1, h2, small, cell = styles()
    pilot = sorted({R.parts[i]["province_name"] for i in range(len(R.parts)) if R.parts[i]["portal"]})
    s = []
    s.append(Paragraph("%s" % P.FOUNDATION_LEGAL, small))
    s.append(Paragraph("%s: Operating Policy" % P.SCHEME, h1))
    s.append(Paragraph("2027 edition. Approved by the Board of Trustees on 12 February 2027 and "
                       "effective from 1 April 2027. Supersedes the 2025 edition.", small))
    s.append(Paragraph("Version 6.0 &bull; Owner: %s &bull; Next review: February 2029" % P.OFFICE, small))
    s.append(Spacer(1, 5))

    s.append(Paragraph("1 Scope", h2))
    s.append(Paragraph(
        "1.1 The %s (the scheme) reimburses accredited community partners for one-off relief awards "
        "made to households in the provinces in which the Foundation works." % P.SCHEME, body))
    s.append(Paragraph(
        "1.2 An award is one household's relief application resolved with a payment. The partner "
        "resolves the application, pays the household and claims the amount back from the "
        "Foundation.", body))
    s.append(Paragraph(
        "1.3 A partner is accredited by the %s for a term of three years and is listed in the "
        "accreditation register with the date from which it may claim." % P.OFFICE, body))

    s.append(Paragraph("4 Claims", h2))
    s.append(Paragraph(
        "4.1 A partner claims monthly. A return states, for each household-size band in Schedule A, "
        "the number of awards the partner resolved with a payment in the claim month and the total "
        "amount of those awards.", body))
    s.append(Paragraph(
        "4.2 The Office may reopen a claim month for reconciliation, in which case returns for that "
        "month are filed as preliminary. A revised return for the same partner and month supersedes "
        "the preliminary return in full.", body))
    s.append(Paragraph(
        "4.3 A recovery of a previously paid award is claimed as a separate negative line in the "
        "month of recovery and is not netted against an award.", body))
    s.append(Paragraph(
        "4.4 A partner in a province named in Schedule C records each award on the partner claim "
        "portal as it is resolved, in addition to the monthly return under 4.1.", body))
    s.append(Paragraph(
        "4.5 An award is claimed on the return for the month in which the partner paid the "
        "household.", body))

    s.append(Paragraph("7 Pre-payment referral", h2))
    s.append(Paragraph(
        "7.1 An award that exceeds the referral ceiling for its household-size band is referred to "
        "the %s before the partner's claim for it is paid." % P.DESK, body))
    s.append(Paragraph(
        "7.2 The Desk reviews a referred award against the household's circumstances and the "
        "receipted costs, and either releases the claim or returns it to the partner with a note.", body))
    s.append(Paragraph(
        "7.3 The ceiling is a review trigger. It is not a limit on the amount of an award and a "
        "partner is not restricted to awards below it.", body))

    s.append(Paragraph("9 The referral ceiling for 2027/28", h2))
    s.append(Paragraph(
        "9.1 For each scheme year the Board sets a referral ceiling at the reference household size "
        "of four persons. The ceiling for the 2027/28 scheme year is the lowest amount, being a "
        "multiple of one hundred pesos, such that when the ceiling for each other band is generated "
        "from it by the factor in Schedule B and rounded up to the nearest one hundred pesos, the "
        "share of the scheme's awards that would be referred under clause 7 is at or below 6.0 per "
        "cent.", body))
    s.append(Paragraph(
        "9.2 For 9.1 the scheme's awards are the awards of all accredited partners at the award "
        "volumes and award levels of the twelve months to 31 March 2027. An award is counted in the "
        "month in which it was paid.", body))
    s.append(Paragraph(
        "9.3 An award is referred if it exceeds the ceiling for its band. Awards are counted gross.", body))
    s.append(Paragraph(
        "9.4 The scheme's awards for 9.1 are every award within those twelve months, including the "
        "awards of a partner accredited or withdrawn during them. A partner accredited after 31 "
        "March 2027 is not within them.", body))
    s.append(Paragraph(
        "9.5 The share in 9.1 is the reviewing capacity the Board has agreed for the Desk for the "
        "scheme year.", body))

    s.append(Paragraph("Schedule A: household-size bands", h2))
    s.append(grid([["Band", "Household size"]] + [[RC.legacy_band(b), P.BAND_LABEL[b]] for b in P.BANDS],
                  [60, 220], cell))
    s.append(Paragraph("The band codes B1 to B6 used in returns to June 2026 correspond to HH1 to "
                       "HH6 in the same order.", small))

    s.append(Paragraph("Schedule B: the equivalence scale and rounding", h2))
    s.append(Paragraph(
        "B.1 The ceiling for a household-size band is generated from the reference-size ceiling by "
        "the factor for that band.", body))
    s.append(Paragraph(
        "B.2 The factor for a band is the square root of the ratio of the band's household size to "
        "the reference household size of four persons, taken to three decimal places. The band's "
        "household size is the lower bound of the band.", body))
    s.append(Paragraph(
        "B.3 A ceiling generated under B.1 is rounded up to the nearest one hundred pesos.", body))
    s.append(Paragraph(
        "B.4 The factor is applied to the reference-size ceiling before the rounding in B.3, and the "
        "rounding is applied once.", body))

    s.append(Paragraph("Schedule C: partner claim portal", h2))
    s.append(Paragraph(
        "The portal has been in use for awards resolved from 1 January 2025 by accredited partners "
        "in %s." % " and ".join(pilot), body))

    s.append(Paragraph("Schedule D: eligible components", h2))
    s.append(grid([["Code", "Component", "Eligible from"]] +
                  [[c, n, "1 July 2024"] for c, n, _ in P.COMPONENTS_STD] +
                  [[c, n, "1 July 2025"] for c, n, _ in P.COMPONENTS_UPPER], [50, 230, 90], cell))

    s.append(Paragraph("Amendment history", h2))
    s.append(grid([["Date", "Amendment"],
                   ["1 July 2025", "Schedule D: components SHL and EQP added to the eligible "
                                   "components by board resolution 2025-14."],
                   ["1 July 2026", "Claim return layout revised by the Office: band column labels "
                                   "and a header block, and the return status column."],
                   ["12 February 2027", "Clause 7 and clause 9 inserted. The pre-payment referral "
                                        "applies from the 2027/28 scheme year."]],
                  [80, 390], cell))
    d = doc(path, "%s Operating Policy 2027" % P.SCHEME, P.OFFICE,
            "%s operating policy, 2027 edition" % P.SCHEME_SHORT)
    foot = footer("%s - %s operating policy, 2027 edition" % (P.FOUNDATION, P.SCHEME_SHORT))
    with stamped(epoch(2027, 2, 16, 11, 20)):
        d.build(s, onFirstPage=foot, onLaterPages=foot, canvasmaker=org_canvas())
    finalise_pdf(path)


# ------------------------------------------------------------------ the published summary (prior year)
def published_summary(path, R, W):
    body, h1, h2, small, cell = styles()
    fy = [t for t in sim.MONTHS if (2025, 7) <= t <= (2026, 6)]
    byprov, bycomp, pilotcell = {}, {}, {}
    pilot = {R.parts[i]["province"] for i in range(len(R.parts)) if R.parts[i]["portal"]}
    for a in W["awards"]:
        if a["month"] not in fy:
            continue
        pt = R.parts[a["partner"]]
        p = byprov.setdefault(pt["province"], [set(), 0, 0.0])
        p[0].add(a["partner"])
        p[1] += 1
        p[2] += a["amount"]
        bycomp[a["component"]] = bycomp.get(a["component"], 0) + 1
        if pt["province"] in pilot:
            k = (a["band"], a["component"])
            pilotcell[k] = pilotcell.get(k, 0) + 1
    s = [Paragraph(P.FOUNDATION_LEGAL, small),
         Paragraph("Annual Report 2025/26: %s, partner summary" % P.SCHEME, h1),
         Paragraph("Extract from the published annual report, laid before the Board of Trustees on "
                   "30 September 2026. Amounts are in thousands of pesos.", small),
         Spacer(1, 5),
         Paragraph("Awards and amounts by province, fiscal year 2025/26", h2)]
    rows = [["Province", "Partners claiming", "Awards", "Amount claimed (PHP '000)",
             "Average award (PHP)"]]
    tp, ta, tv = 0, 0, 0.0
    for code, name, _, _ in P.PROVINCES:
        if code not in byprov:
            continue
        pr, n, v = byprov[code]
        rows.append([name, len(pr), "{:,}".format(n), "{:,.0f}".format(v / 1000.0),
                     "{:,.0f}".format(v / n)])
        tp += len(pr)
        ta += n
        tv += v
    rows.append(["All provinces", tp, "{:,}".format(ta), "{:,.0f}".format(tv / 1000.0),
                 "{:,.0f}".format(tv / ta)])
    s.append(grid(rows, [105, 80, 60, 115, 95], cell, align_right=(1, 2, 3, 4)))
    s.append(Paragraph("Awards by eligible component, all provinces", h2))
    cn = {c: n for c, n, _ in P.COMPONENTS_STD + P.COMPONENTS_UPPER}
    s.append(grid([["Code", "Component", "Awards"]] +
                  [[c, cn[c], "{:,}".format(bycomp.get(c, 0))] for c in cn],
                  [45, 220, 70], cell, align_right=(2,)))
    s.append(Paragraph("Partner claim portal provinces: awards by band and component", h2))
    comps = list(cn)
    rows = [["Band"] + comps]
    for b in P.BANDS:
        r = [RC.legacy_band(b)]
        for c in comps:
            v = pilotcell.get((b, c), 0)
            r.append("n/p" if v < 25 else "{:,}".format(v))
        rows.append(r)
    s.append(grid(rows, [45] + [62] * len(comps), cell, align_right=tuple(range(1, len(comps) + 1))))
    s.append(Paragraph("Cells of fewer than 25 awards are not published separately (n/p) and are "
                       "included in the totals above. Figures are as at the 2025/26 year end and "
                       "are not restated for later revisions to a claim month.", small))
    d = doc(path, "%s annual report 2025/26 partner summary" % P.SCHEME_SHORT, "%s" % P.FOUNDATION,
            "published partner summary, fiscal year 2025/26")
    foot = footer("%s - annual report 2025/26 extract" % P.FOUNDATION)
    with stamped(epoch(2026, 9, 30, 15, 45)):
        d.build(s, onFirstPage=foot, onLaterPages=foot, canvasmaker=org_canvas())
    finalise_pdf(path)


# ------------------------------------------------------------------ the request
def request(path, F):
    text = """From: %s
To: Grants Analytics <grants.analytics@%s>
Cc: %s
Date: Thu, 06 May 2027 08:41:12 +0800
Subject: Referral ceiling schedule for the June board
Message-ID: <20270506084112.2184@%s>
MIME-Version: 1.0
Content-Type: text/plain; charset="utf-8"
Content-Transfer-Encoding: 7bit

Morning,

The trustees take the 2027/28 referral ceiling schedule at the June meeting, so I need the schedule
by 24 May. Clause 9 went into the policy in February and this is the first year we set one, so there
is nothing to roll forward.

Ronald has put the round's folder on the shared drive: the claim returns the claims system holds,
the portal export for the pilot provinces, the accreditation register and the reference tables, and
the 2027 edition of the policy. The returns for November and February were reopened for the
reconciliation exercise, so take care with those two.

What I need is the schedule itself, the share of awards it would refer in each band and the share of
value, and the scale factor you applied, plus a short paper the trustees can read without the
workbook. The Desk's capacity is the constraint the Board signed up to, so the share matters as much
as the amounts.

Thanks,
%s
%s
""" % (addr("assurance"), P.DOMAIN, addr("systems"), P.DOMAIN, P.PEOPLE["assurance"][0],
       P.PEOPLE["assurance"][2])
    open(path, "w", encoding="utf-8", newline="\n").write(text)


# ------------------------------------------------------------------ the data dictionary
def dictionary(path, F):
    D = dict(F)
    D["office"] = P.OFFICE
    D["extract"] = P.EXTRACT_DATE.isoformat()
    t = """# Data dictionary

Household Relief Grant, round folder for the 2027/28 referral ceiling schedule.
Prepared by the %(office)s. Extract date %(extract)s. Issue 3.

A row of a claim return is one partner, one household-size band and one claim month. A row of the
portal export is one award. An award is counted in the month in which the partner paid the
household, and `claim_return_month` on a portal row is that month.

## %(ret2627)s, %(ret2526)s, %(ret2425)s

The monthly claim returns the claims system holds. The 2026/27 workbook carries the layout the
Office introduced on 1 July 2026: a four-row header block above the table, the band column labelled
`household_size_band` with the HH1 to HH6 codes, `awards` and `total_claimed_php` in place of
`award_count` and `claim_total_php`, and a `return_status` column. The two earlier workbooks carry
the former layout and the B1 to B6 band codes.

| field | meaning |
|---|---|
| `return_month`, `claim_month` | the claim month, as `YYYY-MM` in the earlier layout and as `Mon-YYYY` in the 2026/27 layout |
| `partner_code` | the partner's code in the grants ledger |
| `hh_band`, `household_size_band` | the household-size band of the awards on the row |
| `award_count`, `awards` | the number of awards the partner resolved with a payment in the month, in that band |
| `claim_total_php`, `total_claimed_php` | the total amount of those awards, in pesos |
| `return_status` | `FINAL` a return for a month that was not reopened; `PRELIM` a return for a reopened month; `REVISED` a return filed after `PRELIM` for the same partner and month |

A negative `awards` and amount is a recovery of an award paid in an earlier month, claimed as its own
line in the month of recovery.

## %(portal)s

The partner claim portal's export, one row per award, for awards resolved from January 2025 by
partners in the portal provinces.

| field | meaning |
|---|---|
| `portal_award_ref` | the portal's reference for the award |
| `portal_partner_id` | the partner's identifier in the portal |
| `partner_ref_portal` | the partner reference as the portal stores it, loaded from a flat file when the partner was enrolled; `%(xwalk)s` maps it to `partner_code` in the grants ledger |
| `province_code` | the province of the partner |
| `claim_return_month` | the claim month on whose return the award was claimed |
| `application_date`, `decision_date`, `payment_date` | the dates the household applied, the partner resolved the application and the partner paid the household |
| `household_size` | the number of persons in the household |
| `size_band` | the household-size band, B1 to B6 |
| `component_code` | the eligible component of the award, as Schedule D of the policy lists them |
| `award_amount_php` | the amount reimbursed to the partner for the award, in pesos |

## %(register)s

One row per partner the Foundation has accredited. `accreditation_status` is `ACTIVE`, `WITHDRAWN`
where the partner has given up its accreditation, or `LAPSED` where a term ended and was not
renewed; `status_date` carries the date of a withdrawal or a lapse.

## %(xwalk)s

The mapping between portal accounts and the grants ledger. `account_status` is `LIVE` for an account
that records awards, or `OPENED_NOT_LIVE` for an account opened for a partner that did not join the
portal.

## %(bands)s, %(prov)s

Reference tables. `legacy_band_code` carries the HH1 to HH6 codes and `band_code` the B1 to B6 codes
for the same bands. `household_size_to` is blank for the open band. `portal_pilot_from` is the date
the portal opened in a province and is blank where it has not.

## %(treasury)s

The Treasury's fortnightly payment runs. `gross_php` is what was paid out to the partner in the run,
`recoveries_php` what was recovered from it, and `net_php` the amount transferred. A run covers the
awards a partner paid in one half of a claim month and is dated in the following month.

## %(followup)s

The 2026 household follow-up survey, a sample of households whose awards were paid in 2026.
`days_to_payment` is the days from the partner's decision to its payment, and `adequacy_rating` the
household's answer on a five-point scale.

## %(manifest)s

What each file in the folder is, what it covers and what it omits.
""" % D
    open(path, "w", encoding="utf-8", newline="\n").write(t)


# ------------------------------------------------------------------ the export manifest
def manifest(path, F, pack):
    import pandas as pd
    about = {
        "portal": ("partner claim portal export, one row per award",
                   "awards resolved 2025-01-01 to 2027-03-31 by partners in the portal provinces",
                   "partners outside the portal provinces, and awards resolved before the portal opened"),
        "ret2627": ("monthly claim returns, fiscal year 2026/27",
                    "claim months July 2026 to March 2027", "claim months after March 2027"),
        "ret2526": ("monthly claim returns, fiscal year 2025/26",
                    "claim months July 2025 to June 2026", "the return status column, introduced in July 2026"),
        "ret2425": ("monthly claim returns, fiscal year 2024/25",
                    "claim months July 2024 to June 2025", "the return status column, introduced in July 2026"),
        "policy": ("the scheme's operating policy, 2027 edition",
                   "the scheme as it stands from 1 April 2027", "the superseded 2025 edition"),
        "register": ("the accreditation register", "every partner accredited since 2014",
                     "partners that applied for accreditation and were not accredited"),
        "xwalk": ("portal account to grants ledger mapping", "every portal account opened",
                  "partners with no portal account"),
        "bands": ("household-size band reference table", "the bands in Schedule A of the policy", ""),
        "prov": ("province reference table", "the eight provinces the Foundation works in", ""),
        "summary": ("extract from the published annual report 2025/26",
                    "fiscal year 2025/26 as published", "later revisions to a claim month"),
        "treasury": ("Treasury fortnightly payment runs", "runs dated February 2026 to April 2027",
                     "runs before February 2026, held in the Treasury's own archive"),
        "followup": ("household follow-up survey 2026", "a sample of awards paid in 2026",
                     "amounts, which the survey does not collect"),
        "dict": ("the data dictionary for this folder", "every file in the folder", ""),
        "request": ("the request the folder was put together for", "", ""),
        "manifest": ("this manifest", "every file in the folder", ""),
    }
    rows = []
    for k in pack:
        c, cov, om = about[k]
        rows.append({"file_name": F[k], "contents": c, "covers": cov, "omits": om,
                     "produced_by": "%s extract, %s" % (P.OFFICE, P.EXTRACT_DATE.isoformat())
                     if k not in ("policy", "summary", "request") else
                     ("Board of Trustees, approved 2027-02-12" if k == "policy" else
                      ("published annual report 2025/26" if k == "summary" else P.PEOPLE["assurance"][0])),
                     "source": P.FOUNDATION_LEGAL, "extract_date": P.EXTRACT_DATE.isoformat(),
                     "licence": "Internal: Foundation record, released for the schedule review"})
    pd.DataFrame(rows).to_csv(path, index=False, lineterminator="\n")


def write_all(target, F, R, W, pack):
    import os
    p = lambda k: os.path.join(target, F[k])
    claim_returns(p("ret2425"), R, [t for t in sim.MONTHS if (2024, 7) <= t <= (2025, 6)], False, "2024/25")
    claim_returns(p("ret2526"), R, [t for t in sim.MONTHS if (2025, 7) <= t <= (2026, 6)], False, "2025/26",
                  recovery=R.recovery_row())
    claim_returns(p("ret2627"), R, [t for t in sim.MONTHS if (2026, 7) <= t <= (2027, 3)], True, "2026/27")
    policy(p("policy"), R)
    published_summary(p("summary"), R, W)
    request(p("request"), F)
    dictionary(p("dict"), F)
    manifest(p("manifest"), F, pack)
