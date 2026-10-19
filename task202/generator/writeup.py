"""Read the shipped bundle, compute every figure the write-up and the goldens owe, and emit
figures.json and submission.md. Nothing here reads the generator's simulated truth."""
import csv
import json
import os

import numpy as np
from openpyxl import load_workbook

import params as P
import solve

WINDOW = [(2026, m) for m in range(4, 13)] + [(2027, m) for m in range(1, 4)]
MON = {"Jan": 1, "Feb": 2, "Mar": 3, "Apr": 4, "May": 5, "Jun": 6, "Jul": 7, "Aug": 8, "Sep": 9,
       "Oct": 10, "Nov": 11, "Dec": 12}


def _month(lab):
    lab = str(lab).strip()
    if "-" in lab and lab[:3] in MON:
        return (int(lab.split("-")[1]), MON[lab[:3]])
    y, m = lab.split("-")
    return (int(y), int(m))


def _band(lab):
    return int(str(lab).strip()[-1])


def read_returns(paths, by_month=False):
    """Every claim return row in the window, the revised return superseding the preliminary.

    `by_month` keys each cell by its claim month as well, which is the partner-band-month grain
    clause 9.2 now rules out.
    """
    rows = {}
    for path in paths:
        ws = load_workbook(path, read_only=True, data_only=True).active
        grid = [[c for c in r] for r in ws.iter_rows(values_only=True)]
        hdr = next(i for i, r in enumerate(grid)
                   if r and any(str(c).strip() in ("return_month", "claim_month") for c in r if c))
        cols = {str(c).strip(): j for j, c in enumerate(grid[hdr]) if c}
        mc = cols.get("return_month", cols.get("claim_month"))
        bc = cols.get("hh_band", cols.get("household_size_band"))
        nc = cols.get("award_count", cols.get("awards"))
        tc = cols.get("claim_total_php", cols.get("total_claimed_php"))
        sc = cols.get("return_status")
        pc = cols["partner_code"]
        for r in grid[hdr + 1:]:
            if not r or r[mc] is None:
                continue
            mo = _month(r[mc])
            if mo not in WINDOW:
                continue
            n = int(r[nc])
            if n < 0:                                  # a recovery is not an award (clause 9.3)
                continue
            st = str(r[sc]).strip() if sc is not None and r[sc] is not None else "FINAL"
            k = (mo, str(r[pc]).strip(), _band(r[bc]))
            if st == "PRELIM" and k in rows:
                continue
            if k in rows and rows[k][2] == "REVISED" and st != "REVISED":
                continue
            rows[k] = (n, float(r[tc]), st)
    cells = {}
    for (mo, code, b), (n, tot, st) in rows.items():
        k = (code, b, mo) if by_month else (code, b)
        c = cells.setdefault(k, [0, 0.0])
        c[0] += n
        c[1] = round(c[1] + tot, 2)
    return cells


def read_crosswalk(path):
    out = {}
    with open(path, newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            out[r["partner_ref_portal"]] = r["ledger_partner_code"].strip()
    return out


def read_portal(path, xwalk, with_month=False):
    """The portal's awards inside the window, keyed to the grants ledger's partner code."""
    out = []
    with open(path, newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            mo = _month(r["claim_return_month"])
            if mo not in WINDOW:
                continue
            row = (xwalk[r["partner_ref_portal"]], _band(r["size_band"]),
                   float(r["award_amount_php"]))
            out.append(row + (mo,) if with_month else row)
    return out


# ------------------------------------------------------------------ the closed period (the asks)
CHECK_FROM, CHECK_TO = (2024, 7), (2025, 12)
REASONS = ["Receipts incomplete", "Household record incomplete", "Duplicate application",
           "Component not eligible", "Arithmetic error on return", "Not recorded"]


def _scheme_year(mo):
    return mo[0] if mo[1] >= 7 else mo[0] - 1


def _accreditation_year(acc, mo):
    """The year a claim month falls in on the partner's own accreditation anniversary."""
    return ((mo[0] * 12 + mo[1]) - (acc[0] * 12 + acc[1])) // 12


def _quarter(mo):
    return _scheme_year(mo), (mo[1] - 7) % 12 // 3 + 1


def _qlabel(q):
    return "%d/%s Q%d" % (q[0], str(q[0] + 1)[2:], q[1])


def _norm(code):
    """The partner's reference, whichever separator the claims system wrote it with."""
    return str(code).strip().replace("/", "-").replace("_", "-").upper()


def _reason(x):
    """The reason as the list records it. The field was free text until the April 2025 release, so
    the earlier entries are the same five reasons in whatever case and punctuation was typed."""
    t = " ".join(str(x or "").split()).rstrip(".").strip()
    if not t:
        return ""
    for r in REASONS:
        if t.lower() == r.lower():
            return r
    return t


def read_closed_returns(paths):
    """What partners claimed in each closed claim month, and the recoveries claimed back in it."""
    claimed, recov = {}, {}
    for path in paths:
        ws = load_workbook(path, read_only=True, data_only=True).active
        grid = [[c for c in r] for r in ws.iter_rows(values_only=True)]
        hdr = next(i for i, r in enumerate(grid)
                   if r and any(str(c).strip() in ("return_month", "claim_month") for c in r if c))
        cols = {str(c).strip(): j for j, c in enumerate(grid[hdr]) if c}
        mc = cols.get("return_month", cols.get("claim_month"))
        nc = cols.get("award_count", cols.get("awards"))
        tc = cols.get("claim_total_php", cols.get("total_claimed_php"))
        pc = cols["partner_code"]
        for r in grid[hdr + 1:]:
            if not r or r[mc] is None:
                continue
            mo = _month(r[mc])
            if not (CHECK_FROM <= mo <= CHECK_TO):
                continue
            k = (_norm(r[pc]), mo)
            d = recov if int(r[nc]) < 0 else claimed
            d[k] = round(d.get(k, 0.0) + float(r[tc]), 2)
    return claimed, recov


_CLAIMED_N = {}


def _claimed_counts(target, F):
    """The awards each closed claim month was claimed for, from the returns."""
    key = os.path.abspath(target)
    if key in _CLAIMED_N:
        return _CLAIMED_N[key]
    out = {}
    for path in [os.path.join(target, F[k]) for k in ("ret2425", "ret2526")]:
        ws = load_workbook(path, read_only=True, data_only=True).active
        grid = [[c for c in r] for r in ws.iter_rows(values_only=True)]
        hdr = next(i for i, r in enumerate(grid)
                   if r and any(str(c).strip() in ("return_month", "claim_month") for c in r if c))
        cols = {str(c).strip(): j for j, c in enumerate(grid[hdr]) if c}
        mc = cols.get("return_month", cols.get("claim_month"))
        nc = cols.get("award_count", cols.get("awards"))
        pc = cols["partner_code"]
        for r in grid[hdr + 1:]:
            if not r or r[mc] is None:
                continue
            mo = _month(r[mc])
            if not (CHECK_FROM <= mo <= CHECK_TO) or int(r[nc]) < 0:
                continue
            k = (_norm(r[pc]), mo)
            out[k] = out.get(k, 0) + int(r[nc])
    _CLAIMED_N[key] = out
    return out


def read_accreditation(path):
    """Each partner's accreditation date, which is the year the claims system's running totals are
    carried on."""
    out = {}
    with open(path, newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            d = r["accredited_from"].strip()
            if d:
                out[_norm(r["partner_code"])] = (int(d[:4]), int(d[5:7]))
    return out


def read_check_register(path, acc):
    """Each closed claim month's own passed and held amounts, per partner.

    Three readings have to be made before the figures mean anything, and the records force all
    three. The register holds more than one check for a partner claim-month where the partner
    answered the hold, and the later check is the one that stands. The two currency columns state
    the partner's running position rather than the month's own figure: an amount passed above the
    month's own claim is impossible. And the year those totals run on is the partner's own
    accreditation year, not the scheme year the return workbooks label, so the month the series
    restarts in is the partner's accreditation anniversary and differs partner by partner;
    differencing inside that year restores the month's own claim less what was passed as what was
    held, in every partner claim-month, and differencing inside the scheme year does not.
    """
    rows = []
    with open(path, newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            rows.append({"code": _norm(r["claims_partner_ref"]), "month": _month(r["claim_month"]),
                         "checked": r["checked_date"].strip(),
                         "passed": float(r["amount_passed_php"]),
                         "held": float(r["amount_held_php"]),
                         "passed_n": int(r["awards_passed"]),
                         "reason": _reason(r["hold_reason"])})
    live = {}
    for r in rows:
        k = (r["code"], r["month"])
        if k not in live or r["checked"] > live[k]["checked"]:
            live[k] = r
    months = {}
    for (code, mo) in live:
        months.setdefault(code, []).append(mo)
    out = {}
    for code, ms in months.items():
        ms.sort()
        prev = {}
        for mo in ms:
            sy = _accreditation_year(acc[code], mo)
            base = prev.get(sy, (0.0, 0.0))
            r = live[(code, mo)]
            out[(code, mo)] = {"passed": round(r["passed"] - base[0], 2),
                               "held": round(r["held"] - base[1], 2),
                               "passed_n": r["passed_n"], "reason": r["reason"]}
            prev[sy] = (r["passed"], r["held"])
    return out, len(rows)


def read_archive(path, run_year_in=None):
    """The archived payment runs, grouped by the claim month each line covers.

    `run_year_in` keeps only the lines whose run date falls in those calendar years, which is the
    reading that drops the catch-up run; the archive covers claim months, not run dates.
    """
    out = {}
    with open(path, newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            d = r["run_date"].strip()               # dd/mm/yyyy, as the Treasury writes it
            if run_year_in is not None and int(d[6:10]) not in run_year_in:
                continue
            k = (_norm(r["partner_code"]), _month(r["claim_months_covered"]))
            out.setdefault(k, []).append((float(r["gross_php"]), float(r["recoveries_php"]),
                                          float(r["net_php"])))
    return out


def transferred(archive, claimed):
    """What the runs transferred against each closed claim month.

    The archive is an export of payment instructions against run lines, so where one instruction
    settled both of a partner's half-month runs its amounts sit on both lines. The lines of such a
    month state the same amounts and together come to more than the partner claimed for the month,
    which cannot be paid, so they are one instruction and count once. The claim returns settle it on
    their own: the check register is not needed and its own reading does not reach this figure.
    """
    out, repeats, ninstr = {}, 0, {}
    for k, lines in archive.items():
        clm = claimed.get(k)
        if clm is not None and round(sum(g for g, _, _ in lines), 2) > clm + 0.005:
            seen, tot = set(), 0.0
            for ln in lines:
                if ln in seen:
                    continue
                seen.add(ln)
                tot = round(tot + ln[2], 2)
            out[k] = tot
            ninstr[k] = len(seen)
            repeats += 1
        else:
            out[k] = round(sum(n for _, _, n in lines), 2)
            ninstr[k] = len(lines)
    return out, repeats, ninstr


def ask_figures(target, F):
    """The closed-period annex: what the runs transferred, what the checks held and why."""
    acc = read_accreditation(os.path.join(target, F["register"]))
    check, raw_rows = read_check_register(os.path.join(target, F["checks"]), acc)
    archive = read_archive(os.path.join(target, F["archive"]))
    claimed, _ = read_closed_returns([os.path.join(target, F[k]) for k in ("ret2425", "ret2526")])
    trans, repeats, ninstr = transferred(archive, claimed)
    quarters = []
    for mo in sorted({k[1] for k in check}):
        if _quarter(mo) not in quarters:
            quarters.append(_quarter(mo))
    tq = {_qlabel(q): 0.0 for q in quarters}
    hq = {_qlabel(q): 0.0 for q in quarters}
    pq = {_qlabel(q): 0.0 for q in quarters}
    iq = {_qlabel(q): 0 for q in quarters}
    rc = {r: 0 for r in REASONS}
    for k, v in check.items():
        q = _qlabel(_quarter(k[1]))
        hq[q] = round(hq[q] + v["held"], 2)
        pq[q] = round(pq[q] + v["passed"], 2)
        if v["held"] > 0.005:
            rc[v["reason"] if v["reason"] in REASONS else "Not recorded"] += 1
    for k, v in trans.items():
        tq[_qlabel(_quarter(k[1]))] = round(tq[_qlabel(_quarter(k[1]))] + v, 2)
    for k, n in ninstr.items():
        iq[_qlabel(_quarter(k[1]))] += n
    cq = {_qlabel(q): 0.0 for q in quarters}
    for k, v in claimed.items():
        cq[_qlabel(_quarter(k[1]))] = round(cq[_qlabel(_quarter(k[1]))] + v, 2)
    rq = {_qlabel(q): 0.0 for q in quarters}
    for k, lines in archive.items():
        seen, tot = set(), 0.0
        for ln in lines:
            if ln in seen:
                continue
            seen.add(ln)
            tot = round(tot + ln[1], 2)
        v = tot if k in trans and len(set(lines)) < len(lines) else round(
            sum(x[1] for x in lines), 2)
        rq[_qlabel(_quarter(k[1]))] = round(rq[_qlabel(_quarter(k[1]))] + v, 2)
    uq = {q: round(cq[q] - hq[q] - rq[q] - tq[q], 2) for q in cq}
    return {"quarters": [_qlabel(q) for q in quarters],
            "_raw": {"transferred": tq, "held": hq, "claimed": cq, "recovered": rq,
                     "unsettled": uq, "passed": pq},
            "transferred": {k: int(round(v)) for k, v in tq.items()},
            "instructions": dict(iq),
            "passed": {k: int(round(v)) for k, v in pq.items()},
            "held": {k: int(round(v)) for k, v in hq.items()},
            "claimed": {k: int(round(v)) for k, v in cq.items()},
            "recovered": {k: int(round(v)) for k, v in rq.items()},
            "unsettled": {k: int(round(v)) for k, v in uq.items()},
            "holds_by_reason": rc,
            "check_rows": raw_rows, "check_months": len(check),
            "instruction_repeats": repeats, "run_months": len(archive)}


def _raw_register(path):
    """The register as it stands, every row, with no reading applied."""
    rows = []
    with open(path, newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            rows.append({"code": _norm(r["claims_partner_ref"]), "month": _month(r["claim_month"]),
                         "checked": r["checked_date"].strip(),
                         "passed": float(r["amount_passed_php"]),
                         "held": float(r["amount_held_php"]),
                         "reason": r["hold_reason"].strip()})
    return rows


def ask_stops(target, F):
    """Every complete competent pipeline for the three annex asks, at its own value.

    Each stop is a response that did the work and stopped at a different reading. The ladder file is
    the generator's; nothing here ships.
    """
    reg = _raw_register(os.path.join(target, F["checks"]))
    acc = read_accreditation(os.path.join(target, F["register"]))
    check, _ = read_check_register(os.path.join(target, F["checks"]), acc)
    archive = read_archive(os.path.join(target, F["archive"]))
    claimed, _ = read_closed_returns([os.path.join(target, F[k]) for k in ("ret2425", "ret2526")])
    trans, _, ninstr = transferred(archive, claimed)
    quarters = []
    for mo in sorted({k[1] for k in check}):
        if _quarter(mo) not in quarters:
            quarters.append(_quarter(mo))
    QL = [_qlabel(q) for q in quarters]
    blank = lambda: {q: 0.0 for q in QL}

    def byq(d):
        out = blank()
        for k, v in d.items():
            out[_qlabel(_quarter(k[1]))] = round(out[_qlabel(_quarter(k[1]))] + v, 2)
        return {k: int(round(v)) for k, v in out.items()}

    def countq(d):
        """A per-claim-month count summed into the quarter of the claim month."""
        out = {q: 0 for q in QL}
        for k, v in d.items():
            out[_qlabel(_quarter(k[1]))] += int(v)
        return out



    def pair(amount, count):
        """One stop of ask A: the pesos settled and the instructions settled on, per quarter."""
        return dict(list(amount.items()) + [("n " + k, v) for k, v in count.items()])

    def pairB(held, passed):
        """One stop of ask B: the pesos the check passed and the pesos it held back, per quarter."""
        return dict(list(held.items()) + [("passed " + k, v) for k, v in passed.items()])

    # ---- A, what the runs transferred
    A = {}
    nlines = {k: len(v) for k, v in archive.items()}
    A["A1 every run line summed as it stands"] = pair(byq(
        {k: round(sum(n for _, _, n in v), 2) for k, v in archive.items()}), countq(nlines))
    A["A2 the gross column summed, recoveries left in"] = pair(byq(
        {k: round(sum(g for g, _, _ in v), 2) for k, v in archive.items()}), countq(nlines))
    A["A3 the over-stated claim-months dropped"] = pair(byq(
        {k: (0.0 if len(set(v)) < len(v) else round(sum(n for _, _, n in v), 2))
         for k, v in archive.items()}),
        countq({k: (0 if len(set(v)) < len(v) else len(v)) for k, v in archive.items()}))
    # the reading the asking sentence closes: the instruction counted once, but attributed to the
    # quarter the run fell in rather than the quarter of the claim month it settles
    byrun, nrun = {}, {}
    with open(os.path.join(target, F["archive"]), newline="", encoding="utf-8") as fh:
        seen = {}
        for r in csv.DictReader(fh):
            k = (_norm(r["partner_code"]), _month(r["claim_months_covered"]))
            trip = (float(r["gross_php"]), float(r["recoveries_php"]), float(r["net_php"]))
            if k in trans and trip in seen.get(k, set()):
                continue                              # the same instruction a second time
            seen.setdefault(k, set()).add(trip)
            d = r["run_date"].strip()
            q = _qlabel(_quarter((int(d[6:10]), int(d[3:5]))))
            byrun[q] = round(byrun.get(q, 0.0) + trip[2], 2)
            nrun[q] = nrun.get(q, 0) + 1
    A["A4 one instruction once, by the quarter of the run date"] = pair(
        {q: int(round(byrun.get(q, 0.0))) for q in QL}, {q: nrun.get(q, 0) for q in QL})
    A["A5 one instruction once, by the quarter of the claim month (the answer)"] = pair(
        byq(trans), countq(ninstr))

    # ---- B, what the check held back
    B = {}
    b1 = {(r["code"], r["month"]): r["held"] for r in sorted(reg, key=lambda r: r["checked"])}
    b1p = {(r["code"], r["month"]): r["passed"] for r in sorted(reg, key=lambda r: r["checked"])}
    B["B1 the two currency columns summed as they stand"] = pairB(byq(b1), byq(b1p))
    seq = {}
    for r in sorted(reg, key=lambda r: (r["code"], r["month"], r["checked"])):
        seq.setdefault(r["code"], []).append(r)
    flat, flatp, first = {}, {}, {}
    for code, rs in seq.items():
        prev = None
        for r in rs:
            k = (code, r["month"])
            if k not in first:
                first[k] = r
            flat[k] = round(r["held"] - (prev["held"] if prev is not None else 0.0), 2)
            flatp[k] = round(r["passed"] - (prev["passed"] if prev is not None else 0.0), 2)
            prev = r
    B["B2 differenced with no period reset at all"] = pairB(byq(flat), byq(flatp))

    def diff(pick_first, year, fld="held"):
        out, prev = {}, {}
        for code, rs in seq.items():
            for r in rs:
                k = (code, r["month"])
                if pick_first and r is not first[k]:
                    continue
                sy = year(code, r["month"])
                base = prev.get((code, sy), 0.0)
                out[k] = round(r[fld] - base, 2)
                prev[(code, sy)] = r[fld]
        return out

    sy_year = lambda c, mo: _scheme_year(mo)
    ay_year = lambda c, mo: _accreditation_year(acc[c], mo)
    B["B3 the later check, differenced inside the scheme year"] = pairB(
        byq(diff(False, sy_year)), byq(diff(False, sy_year, "passed")))
    B["B4 the first check of each claim-month, accreditation year"] = pairB(
        byq(diff(True, ay_year)), byq(diff(True, ay_year, "passed")))
    B["B5 the later check, differenced inside the accreditation year (the answer)"] = pairB(
        byq({k: v["held"] for k, v in check.items()}),
        byq({k: v["passed"] for k, v in check.items()}))

    # ---- C, the claim-months held, by reason
    C = {}
    raw = {r: 0 for r in REASONS}
    for r in reg:
        if r["held"] > 0.005 and r["reason"] in REASONS:
            raw[r["reason"]] += 1
    C["C1 rows with a held amount, by the reason column"] = dict(raw)
    raw2 = {r: 0 for r in REASONS}
    for r in reg:
        if r["held"] > 0.005:
            raw2[r["reason"] if r["reason"] in REASONS else "Not recorded"] += 1
    C["C2 the same, blanks counted as a reason of their own"] = raw2
    raw3 = {r: 0 for r in REASONS}
    for k, v in check.items():
        if v["held"] > 0.005 and v["reason"] in REASONS:
            raw3[v["reason"]] += 1
    C["C3 each month's own hold, blanks dropped"] = raw3
    raw4 = {r: 0 for r in REASONS}
    for k, v in check.items():
        if v["held"] > 0.005:
            raw4[v["reason"] if v["reason"] in REASONS else "Not recorded"] += 1
    C["C4 each month's own hold, a hold with no reason recorded counted (the answer)"] = raw4
    return {"quarters": QL, "A": A, "B": B, "C": C}


def figures(target, F):
    cells = read_returns([os.path.join(target, F[k]) for k in ("ret2526", "ret2627")])
    xwalk = read_crosswalk(os.path.join(target, F["xwalk"]))
    por = read_portal(os.path.join(target, F["portal"]), xwalk)
    mean = {k: v[1] / v[0] for k, v in cells.items()}
    cnt = {k: v[0] for k, v in cells.items()}
    shape = np.array([a / mean[(c, b)] for c, b, a in por], float)
    shape = shape / shape.mean()
    C = solve.lowest(lambda x: solve.share_cells(x, mean, cnt, shape))
    W = dict(mean=mean, cnt=cnt, shape=shape)
    f = solve.figures_for(C, W)
    s0 = solve.share_cells(C, mean, cnt, shape)
    s1 = solve.share_cells(C - P.CANDIDATE_STEP, mean, cnt, shape)
    default = solve.lowest(lambda x: solve.share_cells(x, mean, cnt, np.array([1.0])))
    fig = {
        "call": int(C),
        "call_php": "PHP %s" % "{:,}".format(int(C)),
        "call_down": int(C - P.CANDIDATE_STEP),
        "round_to": P.ROUND_TO,
        "candidate_step": P.CANDIDATE_STEP,
        "share_pct": round(100 * s0, 2),
        "share_down_pct": round(100 * s1, 2),
        "committed_pct": round(100 * P.COMMITTED_SHARE, 1),
        "awards": int(round(f["awards"])),
        "cells": len(cnt),
        "partners": len({k[0] for k in cnt}),
        "portal_partners": len({c for c, b, a in por}),
        "portal_window_rows": len(por),
        "cv": round(float(shape.std(ddof=1) / shape.mean()), 2),
        "referred_share": round(f["referred_share"], 1),
        "referred_value_share": round(f["referred_value_share"], 1),
        "referred_awards": int(round(f["referred"])),
        "ks_lognormal": None, "ks_gamma": None,
        "default_call": int(default),
        "default_share_pct": round(100 * solve.share_cells(default, mean, cnt, shape), 2),
        "bands": {},
    }
    for b in P.BANDS:
        d = f["bands"][b]
        fig["bands"][str(b)] = {
            "ceiling": int(d["ceiling"]), "factor": round(d["factor"], 3),
            "referred_share": round(d["referred_share"], 1),
            "referred_value_share": round(d["referred_value_share"], 1),
            "awards": int(round(d["awards"])),
        }
    from scipy.stats import norm as _n, gamma as _g
    xs = np.sort(shape)
    n = len(xs)
    ecdf = np.arange(1, n + 1) / n
    crit = 1.36 / np.sqrt(n)
    cv = float(shape.std(ddof=1) / shape.mean())
    v = np.log(1 + cv ** 2)
    kk = 1.0 / cv ** 2
    fig["ks_lognormal"] = int(round(float(np.abs(_n.cdf((np.log(xs) + 0.5 * v) / np.sqrt(v)) - ecdf).max()) / crit))
    fig["ks_gamma"] = int(round(float(np.abs(_g.cdf(xs * kk, kk) - ecdf).max()) / crit))
    fig["asks"] = ask_figures(target, F)
    fig["_asks_raw"] = fig["asks"].pop("_raw")
    fig["_test"] = _test_block(fig, mean, cnt, shape, C)
    return fig


def _test_block(fig, mean, cnt, shape, C):
    """What the goldens need beyond the graded figures: the reference band's own histogram."""
    shp = np.sort(shape)
    hi = 2500.0 * np.ceil(2.7 * fig["bands"][str(P.REF_BAND)]["ceiling"] / 2500.0)
    edges = np.linspace(0.0, hi, 46)
    counts = np.zeros(len(edges) - 1)
    for k in sorted(cnt):
        if k[1] != P.REF_BAND:
            continue
        m = mean[k]
        cdf = np.searchsorted(shp, edges / m, side="right") / len(shp)
        counts += cnt[k] * np.diff(cdf)
        counts[-1] += cnt[k] * (1.0 - cdf[-1])        # the last bin is open, so nothing is dropped
    return {"hist_edges": [float(x) for x in edges], "hist_counts": [float(x) for x in counts],
            "default_call": fig["default_call"], "default_share_pct": fig["default_share_pct"]}


def golden_content(fig):
    return fig


def write_figures(path, fig):
    out = {k: v for k, v in fig.items() if not k.startswith("_")}
    with open(path, "w") as fh:
        json.dump(out, fh, indent=2, sort_keys=True)
        fh.write("\n")


# ------------------------------------------------------------------ submission.md
def submission(fig):
    b = fig["bands"]
    g = lambda k, f: b[str(k)][f]
    lines = []
    A = lines.append
    A("# task202 submission")
    A("")
    A("## Tags")
    A("")
    A("**Domain:** General statistics and data science (exploratory data analysis and distribution "
      "assessment).")
    A("**Analytical objective:** Anomaly detection and diagnostics.")
    A("")
    A("## 1. Final Recommendation")
    A("")
    A("**Set the 2027/28 referral ceiling for a four-person household at PHP %s.**" % "{:,}".format(fig["call"]))
    A("The schedule it generates refers %.2f per cent of the awards the scheme resolved in the twelve "
      "months to end-March 2027, inside clause 9's %.1f per cent, and PHP %s refers %.2f per cent. Not "
      "the distribution of the %d partner-band average awards, which carries no spread within a "
      "partner and files PHP %s, a schedule that would refer %.2f per cent. Not the portal's awards "
      "pooled or rescaled to the portfolio, which do not reproduce the average award the returns state "
      "for each of the other partners. Not a lognormal or a gamma at the portal's coefficient of "
      "variation, which the portal's own fit refuses at %d and %d times the five per cent critical "
      "value." % (fig["share_pct"], fig["committed_pct"], "{:,}".format(fig["call_down"]),
                   fig["share_down_pct"], fig["cells"], "{:,}".format(fig["default_call"]),
                   fig["default_share_pct"], fig["ks_lognormal"], fig["ks_gamma"]))
    A("")
    A("## 2. Critical Components")
    A("")
    A("1. The twelve months to end-March 2027 hold **%s** awards in **%d** partner-band cells across "
      "**%d** accredited partners" % ("{:,}".format(fig["awards"]), fig["cells"], fig["partners"]))
    A("2. The award shape measured on the window's portal claims has a coefficient of variation of "
      "**%.2f**" % fig["cv"])
    A("3. At PHP %s the schedule refers **%.2f per cent** of those awards" %
      ("{:,}".format(fig["call"]), fig["share_pct"]))
    A("4. At PHP %s it refers **%.2f per cent**, above the commitment, so PHP %s is the lowest "
      "qualifying whole thousand" % ("{:,}".format(fig["call_down"]), fig["share_down_pct"],
                                          "{:,}".format(fig["call"])))
    A("")
    A("## 3. Step-by-Step Solution")
    A("")
    A("1. Took the claim returns for April 2026 to March 2027 from hrg_claim_returns_fy2025-26.xlsx and "
      "hrg_claim_returns_fy2026-27.xlsx, the revised return superseding the preliminary one for the two "
      "reopened months as the policy requires, and recoveries excluded because the policy counts awards "
      "gross: %s awards in %d partner-band cells." %
      ("{:,}".format(fig["awards"]), fig["cells"]))
    A("2. Divided each cell's claimed total by its award count: each partner-band's own average award, "
      "the only award-level figure the returns state.")
    A("3. Resolved portal_awards_2025-01_2027-03.csv's partner reference to the grants ledger through "
      "partner_identifier_crosswalk.csv and kept the awards claimed in those twelve months: %s award "
      "rows for %d partners." %
      ("{:,}".format(fig["portal_window_rows"]), fig["portal_partners"]))
    A("4. Expressed each of those awards as a multiple of its own partner-band average award: the "
      "award shape, coefficient of variation %.2f, which a single lognormal and a single gamma at that "
      "coefficient of variation fit at %d and %d times the five per cent Kolmogorov-Smirnov critical "
      "value." % (fig["cv"], fig["ks_lognormal"], fig["ks_gamma"]))
    A("5. Rebuilt each cell's award distribution as the cell's own average award times that shape and "
      "weighted each cell by its award count, which is the policy's share of the scheme's awards.")
    A("6. Generated each band's ceiling from a candidate reference ceiling by schedule B.2's factor in "
      "hrg_operating_policy_2027_edition.pdf, rounded up to the nearest PHP 100 under B.3, and took the lowest candidate whose referred share "
      "is at or below %.1f per cent: PHP %s refers %.2f per cent and PHP %s refers %.2f per cent." %
      (fig["committed_pct"], "{:,}".format(fig["call"]), fig["share_pct"],
       "{:,}".format(fig["call_down"]), fig["share_down_pct"]))
    A("7. Cut instead on the distribution of the %d partner-band averages, the ceiling is PHP %s and "
      "the schedule refers %.2f per cent of the window's awards." %
      (fig["cells"], "{:,}".format(fig["default_call"]), fig["default_share_pct"]))
    A("8. Set the 2027/28 referral ceiling for a four-person household at PHP %s." %
      "{:,}".format(fig["call"]))
    A("")
    A("## 4. Deliverable Answers")
    A("")
    A("### relief_referral_ceilings_2027-28.xlsx")
    A("")
    A("1. Each band's 2027/28 referral ceiling in pesos, its referred share of that band's own awards "
      "and of that band's own award value to one decimal place, and the scale factor applied to three "
      "decimals, in band order:")
    for k in P.BANDS:
        A("   - %s: ceiling PHP %s, refers %.1f per cent of awards and %.1f per cent of award value, "
          "scale factor %.3f" % (P.BAND_LABEL[k], "{:,}".format(g(k, "ceiling")),
                                 g(k, "referred_share"), g(k, "referred_value_share"), g(k, "factor")))
    A("2. All bands together: refers %.1f per cent of awards and %.1f per cent of award value." %
      (fig["referred_share"], fig["referred_value_share"]))
    A("")
    A("### referral_ceiling_board_paper.pdf")
    A("")
    A("1. The recommendation on page 1: set the 2027/28 referral ceiling for a four-person household "
      "at PHP %s." % "{:,}".format(fig["call"]))
    A("2. The same six ceilings, referred shares of awards and of award value, and scale factors as "
      "the workbook, in the schedule table on page 1.")
    A("3. The award distribution and the ceilings drawn across household sizes: a histogram of the "
      "four-person band's award amounts with the PHP %s ceiling marked and the referred tail shaded, "
      "beside the six ceilings stepped across the household-size bands with each band's referred share "
      "labelled, on page 1." % "{:,}".format(fig["call"]))
    A("")
    a = fig["asks"]
    A("### claims_annex.xlsx")
    A("")
    A("1. Of each closed scheme quarter's claims, the pesos the Treasury settled and the payment "
      "instructions it settled them on:")
    for q in a["quarters"]:
        A("   - %s: PHP %s on %s instructions"
          % (q, "{:,}".format(a["transferred"][q]), "{:,}".format(a["instructions"][q])))
    A("2. The pesos the pre-payment claim check passed and the pesos it held back in each closed "
      "scheme quarter, from each month's own position:")
    for q in a["quarters"]:
        A("   - %s: passed PHP %s, held PHP %s"
          % (q, "{:,}".format(a["passed"][q]), "{:,}".format(a["held"][q])))
    A("3. The partner claim-months the check held an amount back in, by reason recorded:")
    for r in REASONS:
        A("   - %s: %s" % (r, "{:,}".format(a["holds_by_reason"][r])))
    A("")
    return "\n".join(lines)
