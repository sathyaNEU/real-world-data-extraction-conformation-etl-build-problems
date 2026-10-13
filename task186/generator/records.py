"""The pack's data files, each as its source system exports it."""
import csv
import datetime as dt
import json

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font

import market as M
import survey as SV
import world as W

LOG = "price_change_log_regular_2025-01-01_2026-11-15.csv"
RACK = "terminal_rack_prices_2024-12_2026-11.csv"
REGISTER = "station_register_2026-11-16.xlsx"
CROSSWALK = "station_renumbering_2026-06-08.csv"
SURVEY = "fuel_sales_survey_returns_2026.csv"
FRAME = "fuel_sales_survey_frame_2026-11.csv"
SUBMISSIONS = "price_submissions_2026-11.csv"
BRANDS = "monitored_brand_register.json"
OPSLOG = "reporting_platform_operations_log.xlsx"
BULLETIN = "weekly_price_bulletin_2026.xlsx"
COMPLAINTS = "price_complaints_register_2026.csv"
TERMINAL = "Port Ostley Terminal"
BOLD = Font(bold=True)


def iso(t):
    return t.strftime("%Y-%m-%dT%H:%M:%SZ")


def cents(tenths):
    return "%.1f" % (tenths / 10.0)


def write_log(path, rows):
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["station_number", "effective_utc", "price_cents_per_litre", "received_utc"])
        for sid, eff, price, rec in rows:
            w.writerow([sid, iso(eff), cents(price), iso(rec)])
    return len(rows)


def write_rack(path, rack):
    rows = 0
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["price_date", "terminal", "product", "rack_cad_per_litre"])
        for d in sorted(rack["RUL"]):
            for prod in ("RUL", "PUL", "ULSD"):
                w.writerow([d.isoformat(), TERMINAL, prod, "%.4f" % rack[prod][d]])
                rows += 1
    return rows


def register_number(s):
    return s.new_id or s.old_id


def write_register(path, sim):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stations"
    ws.append(["station_number", "trading_name", "street_address", "area", "site_ownership", "opened", "closed"])
    for c in ws[1]:
        c.font = BOLD
    for s in sorted(sim["stations"], key=lambda s: register_number(s)):
        last = s.closed - dt.timedelta(days=1) if s.closed else W.AS_OF
        ws.append([register_number(s), M.trading_name(s, last), s.street, W.AREA_NAME[s.area],
                   "Company" if s.ownership == "company" else "Dealer",
                   s.opened if s.opened >= W.LOG_START else s.hist_opened, s.closed])
    for col, wd in zip("ABCDEFG", (15, 40, 26, 17, 15, 12, 12)):
        ws.column_dimensions[col].width = wd
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=7):
        for c in row:
            c.number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    hb = wb.create_sheet("Brand history")
    hb.append(["station_number", "brand_code", "valid_from", "valid_to"])
    for c in hb[1]:
        c.font = BOLD
    for s in sorted(sim["stations"], key=lambda s: register_number(s)):
        for b, f, to in s.brands:
            end = to or s.closed                # a closed station's brand ends the day before it closed
            hb.append([register_number(s), b, f, (end - dt.timedelta(days=1)) if end else None])
    for row in hb.iter_rows(min_row=2, min_col=3, max_col=4):
        for c in row:
            c.number_format = "yyyy-mm-dd"
    for col, wd in zip("ABCD", (15, 12, 12, 12)):
        hb.column_dimensions[col].width = wd
    hb.freeze_panes = "A2"
    ab = wb.create_sheet("About")
    for line in ("Station register, " + W.BUREAU + ", " + W.DIVISION,
                 "Exported 16 November 2026 from the price reporting platform's station table.",
                 "Station numbers: area-prefixed numbers from 8 June 2026; stations that closed before then keep "
                 "their five-digit number.",
                 "Brand codes as on the monitored brand register; IND is an unbranded station.",
                 "site_ownership: Company where the brand's owner owns the site, Dealer where the operator owns it.",
                 "opened: the first day the station traded; closed: the first day it no longer traded.",
                 "valid_to is the last day under the brand; blank means current."):
        ab.append([line])
    ab.column_dimensions["A"].width = 110
    wb.save(path)


def write_crosswalk(path, sim):
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["old_number", "new_number"])
        rows = sorted((s.old_id, s.new_id) for s in sim["stations"] if s.old_id and s.new_id)
        w.writerows(rows)
    return len(rows)


def write_survey(path, rets):
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["survey_reference", "reference_month", "processing_month", "revision", "reporting_name",
                    "site_address", "regular_litres", "premium_litres", "diesel_litres", "received_date"])
        for r in rets:
            w.writerow([r["sfn"], r["ref"].strftime("%Y-%m"), r["proc"].strftime("%Y-%m"), r["rev"], r["name"],
                        r["addr"], r["reg"], r["prem"], r["dsl"], r["received"].isoformat()])
    return len(rets)


def write_frame(path, fr):
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["survey_reference", "station_number", "frame_status", "on_frame_since", "closed_on"])
        for r in fr:
            w.writerow([r["sfn"], r["station"], r["status"], r["since"].isoformat(),
                        r["closed"].isoformat() if r["closed"] else ""])
    return len(fr)


def write_submissions(path, subs):
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["submission_id", "station_number", "grade", "effective_utc", "price_cents_per_litre", "channel",
                    "received_utc"])
        for r in subs:
            w.writerow([r["id"], r["station"], r["grade"], iso(r["eff"]), cents(r["price"]), r["channel"],
                        iso(r["received"])])
    return len(subs)


def write_brands(path):
    doc = {"register": "Monitored brands, " + W.ACT, "kept_by": W.BUREAU + ", " + W.DIVISION,
           "as_of": W.AS_OF.isoformat(),
           "brands": [{"brand_code": c, "brand": n, "registered_owner": o, "head_office": h,
                       "monitored_since": "2019-07-01" if c not in ("STL", "GAR") else "2021-04-01"}
                      for c, n, o, h in W.BRANDS]}
    with open(path, "w") as fh:
        json.dump(doc, fh, indent=1)
        fh.write("\n")


def write_opslog(path, stats, sim):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Events"
    ws["A1"] = "Price reporting platform: operations log"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Platform events that delayed or repeated submissions; it records the platform, not stations' pricing. "
                "Version 4.2, refreshed 16 November 2026, owner Platform Operations, " + W.DIVISION + ".")
    ws.append([])
    ws.append(["event", "start_utc", "end_utc", "event_type", "stations_affected", "changes_received_late",
               "changes_received_twice", "note"])
    for c in ws[4]:
        c.font = BOLD
    events = []
    for i, (d, a, b) in enumerate(M.OUTAGES):
        late, twice, sids = stats[i]
        events.append((W.to_utc(d, a), W.to_utc(d, b), "intake outage", len(sids), late, twice,
                       "queued submissions released at recovery; part of the queue was released twice"))
    events.append((W.to_utc(W.RENUMBER, 0), W.to_utc(W.RENUMBER, 0), "station renumbering",
                   sum(1 for s in sim["stations"] if s.old_id and s.new_id), None, None,
                   "area-prefixed station numbers in force; crosswalk issued"))
    for i, e in enumerate(sorted(events, key=lambda e: e[0])):
        ws.append(["OPS-%s-%02d" % (e[0].strftime("%y%m"), i + 1), iso(e[0]), iso(e[1])] + list(e[2:]))
    for col, wd in zip("ABCDEFGH", (14, 22, 22, 20, 17, 21, 22, 70)):
        ws.column_dimensions[col].width = wd
    wb.save(path)



def write_bulletin(path, sim):
    """the bureau's weekly bulletin: the mean posted regular price at 09:00 local each Monday, by area"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Weekly prices"
    ws["A1"] = "Weekly fuel price bulletin 2026: regular gasoline, cents a litre"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Mean of posted prices in force at 09:00 on Mondays, all reporting stations in each area."
    ws.append([])
    header = ["week_of"] + [W.AREA_NAME[a] for a in W.AREA_CODES] + ["Province"]
    ws.append(header)
    for c in ws[4]:
        c.font = BOLD
    mondays = [d for d in W.DAYS if d.year == 2026 and d.weekday() == 0]
    for d in mondays:
        k = W.DAY_INDEX[d]
        vals, allv = [], []
        for a in W.AREA_CODES:
            ps = []
            for s in sim["stations"]:
                if s.area != a or not s.open_on(d):
                    continue
                p = price_at(sim["events"][s.uid], k, 9 * 60)
                if p is not None:
                    ps.append(p / 10.0)
            vals.append(round(float(np.mean(ps)), 1) if ps else None)
            allv += ps
        ws.append([d] + vals + [round(float(np.mean(allv)), 1)])
    for row in ws.iter_rows(min_row=5, max_col=1):
        for c in row:
            c.number_format = "yyyy-mm-dd"
    ws.column_dimensions["A"].width = 12
    ws.freeze_panes = "B5"
    wb.save(path)
    return len(mondays)


def price_at(events, k, minute):
    p = None
    for kk, mm, pr in events:
        if (kk, mm) <= (k, minute):
            p = pr
        else:
            break
    return p


CATEGORIES = ["price rise", "difference between nearby stations", "sign price not the pump price",
              "price changed while queuing", "other"]


def write_complaints(path, sim):
    g = np.random.default_rng([W.SEED, 800])
    rows = []
    seq = 26000
    by_area = {a: [s for s in sim["stations"] if s.area == a] for a in W.AREA_CODES}
    for d in [x for x in W.DAYS if x.year == 2026]:
        n = int(g.poisson(0.9 if d.month != 1 else 1.4))
        for _ in range(n):
            a = str(g.choice(W.AREA_CODES))
            c = [s for s in by_area[a] if s.open_on(d)]
            s = c[int(g.integers(0, len(c)))]
            seq += int(g.integers(1, 6))
            minute = int(g.integers(7 * 60, 21 * 60))
            rows.append(("PC-%d" % seq, (dt.datetime.combine(d, dt.time()) + dt.timedelta(minutes=minute)),
                         str(g.choice(["phone", "web form", "email"], p=[0.45, 0.4, 0.15])), W.AREA_NAME[a],
                         str(g.choice(CATEGORIES, p=[0.38, 0.22, 0.12, 0.13, 0.15]))))
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["complaint_ref", "received_local", "channel", "area", "category"])
        for r in rows:
            w.writerow([r[0], r[1].strftime("%Y-%m-%d %H:%M"), r[2], r[3], r[4]])
    return len(rows)
