"""Assertions on the built pack and goldens: the input gates; the three rules executed on the generator's clean truth
(R1 the raw price record ships and no column holds a margin, a change or a finding; R2 each Layer 1 item is neutral for
the call; R3a the easy route, each station's change against its own day-to-day variation with every filed rule applied,
files Harbourline; R3b the decisive step alone files Tallis; R3c no plant); the ladder's rungs by name with their
margins; every null that holds the false discovery rate files the same stations; the semantic forks SF1 to SF15; the
rounding bins; the asks' stops outside the golden in every cell they touch; the separation of the asks from the call;
the documents' silence; the realism sweep; and the goldens against figures.json."""
import collections
import csv
import datetime as dt
import json
import os
import re

import numpy as np
import openpyxl
from scipy import stats

import analysis as AN
import build as B
import deliver as DL
import docs as D
import market as M
import records as R
import survey as SV
import verify as V
import world as W

N = [0]
HERE = os.path.dirname(os.path.abspath(__file__))
CH = M.CH


def ok(cond, what):
    N[0] += 1
    if not cond:
        raise AssertionError(what)


def _pdf_text(path):
    from pypdf import PdfReader
    return re.sub(r"\s+", " ", " ".join(pg.extract_text() for pg in PdfReader(path).pages))


def bins(sh):
    return {b: DL.tens(sh[b]) for b in W.CODES}


# ------------------------------------------------------------------------------------------------- the procedure
def procedure(margins, brand_of, null="placebo", sel=None):
    """the golden procedure on any set of daily margins: examined stations, each change against the station's own
    changes at dates whose windows hold no levy change, Benjamini-Hochberg at 5 per cent, shares and the inquiry"""
    ex = []
    for uid in sorted(margins):
        b = brand_of(uid)
        if b not in W.CODES:
            continue
        ks, ch = V.window_changes(margins[uid])
        d = ch[ks == CH][0]
        if np.isnan(d):
            continue
        keep = ((ks <= CH - W.WINDOW) | (ks >= CH + W.WINDOW)) & ~np.isnan(ch)
        if sel is not None:
            keep &= sel(ks)
        ref = ch[keep]
        if null == "placebo":
            p = (1 + np.sum(ref >= d)) / (len(ref) + 1)
        elif null == "placebo_two_sided":
            c = np.median(ref)
            p = (1 + np.sum(np.abs(ref - c) >= abs(d - c))) / (len(ref) + 1)
        elif null == "placebo_normal":
            p = 1 - stats.norm.cdf((d - ref.mean()) / ref.std(ddof=1))
        ex.append((uid, b, float(d), float(p)))
    found = AN.bh(np.array([x[3] for x in ex]))
    sh = AN.shares([(u, b, d) for u, b, d, _ in ex], found)
    return {"ex": ex, "found": found, "shares": sh, "call": AN.inquiry(sh)}


def cross_section(ex, centre, scale, rel=None):
    d = np.array([x[2] for x in ex]) if rel is None else rel
    p = 1 - stats.norm.cdf((d - centre(d)) / scale(d))
    found = AN.bh(p)
    return AN.shares([(u, b, dd) for u, b, dd in ex], found)


def iqr(d):
    return (np.percentile(d, 75) - np.percentile(d, 25)) / 1.349


def mad(d):
    return 1.4826 * np.median(np.abs(d - np.median(d)))


# ------------------------------------------------------------------------------------------------- daily price variants
def daily_variant(events, closed, kind):
    """tenths of a cent: noon, the mean of the prices posted in the day (the price in force when none), the 24-hour
    time-weighted average"""
    out = np.full(M.N, np.nan)
    ev = sorted(events)
    j, cur = 0, None
    for k in range(ev[0][0] if ev else M.N, M.N):
        if closed is not None and W.DAYS[k] >= closed:
            break
        today = []
        while j < len(ev) and ev[j][0] == k:
            today.append((ev[j][1], ev[j][2]))
            j += 1
        if kind == "noon":
            p = cur
            for minute, q in today:
                if minute <= 720:
                    p = q
            out[k] = np.nan if p is None else p
        elif kind == "mean_posted":
            ps = [q for _, q in today] or ([cur] if cur is not None else [])
            out[k] = np.mean(ps) if ps else np.nan
        elif kind == "tw24":
            if cur is not None:
                tot, t, p = 0.0, 0, cur
                for minute, q in today:
                    tot += p * (minute - t)
                    t, p = minute, q
                out[k] = (tot + p * (1440 - t)) / 1440.0
        if today:
            cur = today[-1][1]
    return out


def utc_events(events):
    out = []
    for k, minute, p in events:
        t = W.to_utc(W.DAYS[k], minute)
        kk = W.DAY_INDEX.get(t.date())
        if kk is not None:
            out.append((kk, t.hour * 60 + t.minute, p))
    return sorted(out)


# ------------------------------------------------------------------------------------------------- the asks' stops
def litres_stop(sim, rows, month_key, station, brand_rule):
    out = collections.defaultdict(int)
    for r in rows:
        m = r[month_key]
        if m not in SV.ASK_MONTHS:
            continue
        s = station(r)
        b = brand_rule(s, m)
        if b in W.CODES:
            out[(b, m)] += r["reg"]
    return {b: [SV.thousands(out[(b, m)]) for m in SV.ASK_MONTHS] for b in W.CODES}


def latest_returns(rets):
    latest = {}
    for r in rets:
        k = (r["sfn"], r["ref"])
        if k not in latest or r["rev"] > latest[k]["rev"]:
            latest[k] = r
    return list(latest.values())


def priced_variant(sim, subs, grades=("RUL",), channel="API", need=3, share=0.5, by="push"):
    """stations priced by the brand on 1 December 2026 from the November submissions: by the brand's push times (the
    minutes of the day at which `need` or more of its stations submit through `channel`), or by shared minutes (the
    same minute stamp carrying `need` or more of the brand's submissions)"""
    st = sim["by_uid"]
    pushers = collections.defaultdict(set)
    per = collections.defaultdict(list)
    for r in subs:
        if r["grade"] not in grades:
            continue
        s = st[r["uid"]]
        b = s.brand_on(r["eff"].date())
        key = (b, r["eff"].hour * 60 + r["eff"].minute) if by == "push" else (b, r["eff"])
        if channel is None or r["channel"] == channel:
            pushers[key].add(r["uid"] if by == "push" else (r["uid"], r["grade"], r["id"]))
        per[r["uid"]].append(key)
    hot = {k for k, v in pushers.items() if len(v) >= need}
    out = {b: 0 for b in W.CODES}
    for uid, keys in per.items():
        s = st[uid]
        b = s.brand_on(W.SCOPE_DATE)
        if b in W.CODES and s.open_on(W.SCOPE_DATE) and sum(1 for k in keys if k in hot) / len(keys) >= share:
            out[b] += 1
    return out


def run(target, golden, sim, meta_path):
    mc, fig = sim["mc"], sim["fig"]
    st = sim["by_uid"]
    figj = json.load(open(os.path.join(HERE, "figures.json")))
    ok(figj == json.loads(json.dumps(fig)), "figures.json is this build's figures")
    files = sorted(os.listdir(target))
    texts = {f: _pdf_text(os.path.join(target, f)) for f in files if f.endswith(".pdf")}
    texts.update({f: open(os.path.join(target, f), encoding="utf-8").read() for f in files if f.endswith((".md", ".json"))})
    standard, schedule, notes = texts[D.STANDARD], texts[D.SCHEDULE], texts[D.NOTES]
    shares = mc["shares"]
    gold_bins = bins(shares)
    vsrc = open(os.path.join(HERE, "verify.py")).read()
    main_src = vsrc.split("def main_call")[1].split("# ------")[0]
    ask_src = vsrc.split("def litres")[1].split("def tens")[0]

    # ---- input gates
    meta = json.load(open(meta_path))
    ok(len(files) == len(B.FILES) >= 10, "ten or more files")
    ok(len({f.rsplit(".", 1)[1] for f in files}) >= 3, "three or more formats")
    ok(meta["input_gates"]["largest_file_rows"] >= 25000, "a file of 25,000 rows or more")
    ok(len(meta["distractor_files"]) >= 2 and all(f in files for f in meta["distractor_files"]), "two distractors")
    for f in meta["distractor_files"]:
        ok("distract" not in f, "distractor not named as such")
        ok(f not in vsrc, "the solution never reads " + f)
    ok("weekly_price_bulletin" not in vsrc and "price_complaints" not in vsrc, "the verifier reads neither distractor")
    for prefix in ("fuel_sales_survey", "price_submissions"):
        ok(prefix not in main_src, "separation: the main call's path never reads " + prefix)
    for prefix in ("price_change_log", "terminal_rack", "carbon_levy"):
        ok(prefix not in ask_src, "separation: the asks never read " + prefix)
    ok("site_ownership" not in vsrc.split("def litres")[0], "separation: the main call never reads the ownership column")

    # ---- R1: the field's raw material ships, the deciding quantity does not
    with open(os.path.join(target, R.LOG), encoding="utf-8") as fh:
        log = list(csv.DictReader(fh))
    ok(len(log) == len(sim["log"]) > 150000, "R1 the minute log of every regular price change ships (%d rows)" % len(log))
    ok(list(log[0]) == ["station_number", "effective_utc", "price_cents_per_litre", "received_utc"],
       "R1 the log carries the change as submitted and nothing derived")
    for f in files:
        if f.endswith(".csv"):
            head = open(os.path.join(target, f), encoding="utf-8").readline().lower()
            for w in ("margin", "change", "finding", "found", "excess", "pass", "p_value", "null", "cycle", "flag",
                      "priced_by", "price_setter"):
                ok(w not in head, "R1 no shipped column holds a derived quantity (%s in %s)" % (w, f))
    for f, t in texts.items():
        if f != R.BRANDS:
            for nm in W.NAME.values():
                ok(nm not in t, "R1 %s names no brand (%s)" % (f, nm))
    ok("price_change_log" in main_src and "window_changes" in main_src and "ref >= d" in main_src,
       "R1 the verifier's decisive computation reads the log and each station's own changes")

    # ---- R2: each Layer 1 item is real and neutral
    def same_call(res, what):
        ok(res["call"] == mc["call"] and bins(res["shares"]) == gold_bins,
           "%s files the golden (%s, %s)" % (what, res["call"], {b: res["shares"][b] for b in W.CODES}))
    brand_change = lambda uid: st[uid].brand_on(W.CHANGE)
    rk = sim["rack_c"]
    gold = procedure(sim["margin"], brand_change)
    ok(gold["call"] == mc["call"] == "TAL" and gold["shares"] == shares, "the checks' procedure reproduces the golden")
    # UTC days for the local window (SF13)
    ut = {s.uid: AN.margins(AN.daily_prices(utc_events(sim["events"][s.uid]), closed=s.closed), rk)
          for s in sim["stations"]}
    same_call(procedure(ut, brand_change), "R2 SF13 the daily window read in UTC")
    # changes received late and twice after outages (SF14): exact repeats, and the received time for the effective
    keyed = collections.Counter((r[0], r[1]) for r in sim["log"])
    prices = collections.defaultdict(set)
    for r in sim["log"]:
        prices[(r[0], r[1])].add(r[2])
    ok(all(len(v) == 1 for v in prices.values()), "R2 SF14 a change received twice is an exact repeat")
    rep = sum(1 for v in keyed.values() if v > 1)
    ok(rep >= 30, "R2 the outages re-sent %d changes" % rep)
    ok(all(not (W.CHANGE - dt.timedelta(days=28) <= d < W.CHANGE + dt.timedelta(days=28)) for d, a, b in M.OUTAGES),
       "R2 no outage in the change windows")
    uid_of = {}
    for s in sim["stations"]:
        for x in (s.old_id, s.new_id):
            if x:
                uid_of[x] = s.uid
    rec = collections.defaultdict(list)
    for sid, eff, price, r in sim["log"]:
        t = V.to_local(r.strftime("%Y-%m-%dT%H:%M:%SZ"))
        k = W.DAY_INDEX.get(t.date())
        if k is not None:
            rec[uid_of[sid]].append((k, t.hour * 60 + t.minute, price))
    rv = {s.uid: AN.margins(AN.daily_prices(sorted(rec[s.uid]), closed=s.closed), rk) for s in sim["stations"]}
    same_call(procedure(rv, brand_change), "R2 SF14 the received time read as the effective time, repeats kept")
    # the renumbering (SF15): the two numbers read as two stations
    two = {}
    for s in sim["stations"]:
        m = sim["margin"][s.uid].copy()
        if s.old_id and s.new_id:
            m[W.DAY_INDEX[W.RENUMBER]:] = np.nan
        two[s.uid] = m
    same_call(procedure(two, brand_change), "R2 SF15 old numbers read as stations of their own")
    ok(all(not s.new_id or s.old_id for s in sim["stations"] if s.opened < W.CHANGE - dt.timedelta(days=28)),
       "R2 SF15 every station examined carries its old number across the change windows")
    # the rack on days the terminal does not publish (SF12): interpolated instead of carried
    pub = sorted(sim["rack"]["RUL"])
    xs = np.array([(d - W.LOG_START).days for d in pub], float)
    ys = np.array([sim["rack"]["RUL"][d] * 100 for d in pub])
    interp = np.interp(np.arange(M.N, dtype=float), xs, ys)
    ok(np.max(np.abs(interp - rk)) > 0.2, "SF12 interpolation moves the rack on unpublished days")
    iv = {u: sim["margin"][u] + rk - interp for u in sim["margin"]}
    same_call(procedure(iv, brand_change), "R2 SF12 the rack interpolated over unpublished days")
    # the brand history (SF8): today's brand for the brand at the change
    same_call(procedure(sim["margin"], lambda u: st[u].brands[-1][0]), "R2 SF8 today's brand")
    moved = [s for s in sim["stations"] if s.brands[-1][0] != s.brand_on(W.CHANGE) and s.brand_on(W.CHANGE)]
    ok(len(moved) >= 4, "R2 four or more stations changed brand after the change")
    for s in sim["stations"]:
        if len(s.brands) > 1:
            ok(s.brands[1][1] <= W.CHANGE - dt.timedelta(days=28) or s.brands[1][1] >= W.CHANGE + dt.timedelta(days=28),
               "R2 no brand change inside the change windows")

    # ---- R3a, R3b: the easy route and the decisive step, on the clean truth
    nv = mc["naive"]
    lead = lambda sh: sorted(W.CODES, key=lambda b: -sh[b][0] / sh[b][1])
    rate = lambda sh, b: sh[b][0] / sh[b][1]
    ok(AN.inquiry(nv["welch"]) == "HBL", "R3a the easy route (Welch per station, BH 5 per cent) files Harbourline")
    w = nv["welch"]
    ok(rate(w, "HBL") / rate(w, lead(w)[1]) >= 1.2, "R3a Harbourline leads the easy route by 1.2 times")
    ok(lead(w).index("TAL") >= 3, "R3a Tallis fourth or lower on the easy route (%d)" % (lead(w).index("TAL") + 1))
    ok(all(DL.tens(w[b]) != gold_bins[b] for b in W.CODES), "R3a every share off its bin on the easy route")
    for k in ("mann_whitney", "newey_west_3", "newey_west_7", "newey_west_14", "province_relative"):
        sh = nv[k]
        ok(AN.inquiry(sh) == "HBL", "R3b %s alone still files Harbourline" % k)
        ok(rate(sh, "HBL") / rate(sh, lead(sh)[1]) >= 1.1, "R3b %s Harbourline clear by 1.1 times" % k)
        ok(rate(sh, "HBL") >= 1.2 * rate(sh, "TAL"), "R3b %s Tallis behind Harbourline by 1.2 times or more" % k)
    ok(sum(1 for k in ("welch", "mann_whitney", "newey_west_3", "newey_west_7", "newey_west_14", "province_relative")
           if lead(nv[k])[1] == "TAL") <= 3, "Tallis second on at most three within-window variants")
    ok(AN.inquiry(shares) == "TAL", "R3b the decisive step alone (the null from the station's own changes) files Tallis")
    mean = mc["mean_change"]
    ok(max(W.CODES, key=lambda b: mean[b]) == "HBL" and mean["HBL"] / sorted(mean.values())[-2] >= 1.15,
       "R0 each brand's mean margin change: Harbourline largest by 1.15 times")
    # R3c: no plant: no repeated block in any station's daily price series, the cycles are every area's
    for s in sim["stations"]:
        dp = np.round(sim["daily"][s.uid], 2)
        seen, reps = set(), 0
        for i in range(0, M.N - 14):
            blk = dp[i:i + 14]
            if np.isnan(blk).any() or len(set(blk)) < 4:
                continue
            key = tuple(blk)
            reps += key in seen
            seen.add(key)
        ok(reps == 0, "R3c no repeated fortnight in station %d's daily prices" % s.uid)
    periods = [sim["params"][a][0] for a in W.AREA_CODES]
    ok(min(periods) >= 35 and max(periods) <= 63, "R3c every area cycles over five to nine weeks")
    for f, t in texts.items():
        low = t.lower()
        for wd in ("cycle", "restor", "undercut", "placebo", "empirical", "day-to-day", "t-test", "welch", "autocorrel",
                   "phase", "seasonal", "same-area", "competitor", "reissu", "reuse", "freed", "commission agent",
                   "ownership", "dealer group", "vendor"):
            ok(wd not in low, "R3c and silence: %s says nothing of %r" % (f, wd))
        ok("\u2014" not in t and "\u2013" not in t, "no dash in " + f)

    # ---- every null that holds the rate files the same call and share bins; the routes that cannot find anything
    # find nothing
    VALID = {}
    ex3 = [(u, b, d) for u, b, d, _ in gold["ex"]]
    for k, v in mc["valid"].items():
        ok(AN.inquiry(v) == "TAL" and bins(v) == gold_bins, "SF7 the %s null files the golden (%s)" % (k, v["TAL"]))
        VALID[k] = v["TAL"][0]
    variants = {
        "placebo, dates before the change only": procedure(sim["margin"], brand_change, sel=lambda ks: ks <= CH - 28),
        "placebo, dates after the change only": procedure(sim["margin"], brand_change, sel=lambda ks: ks >= CH + 28),
        "placebo, normal approximation": procedure(sim["margin"], brand_change, null="placebo_normal"),
        "placebo, two-sided": procedure(sim["margin"], brand_change, null="placebo_two_sided"),
    }
    for k, v in variants.items():
        ok(v["call"] == "TAL" and bins(v["shares"]) == gold_bins, "SF7 %s files the golden" % k)
        VALID[k] = v["shares"]["TAL"][0]
    for nm, sc in (("interquartile range", iqr), ("MAD", mad)):
        v = cross_section(ex3, np.median, sc)
        ok(AN.inquiry(v) == "TAL" and bins(v) == gold_bins, "SF7 empirical null, median and %s" % nm)
        VALID["empirical, " + nm] = v["TAL"][0]
    dd = {u: d for u, b, d in ex3}
    rel = np.array([d - np.mean([dd[u2] for u2, b2, _ in ex3 if st[u2].area == st[u].area and u2 != u and b2 != b])
                    for u, b, d in ex3])
    for nm, sc in (("MAD", mad), ("IQR", iqr), ("SD", lambda x: x.std(ddof=1))):
        v = cross_section(ex3, np.median, sc, rel=rel)
        ok(AN.inquiry(v) == "TAL" and bins(v) == gold_bins, "SF7 same-area competitors, %s scale" % nm)
        VALID["same-area, " + nm] = v["TAL"][0]
    area_mean = {}
    for a in W.AREA_CODES:
        rows = [sim["margin"][u] for u, b, d in ex3 if st[u].area == a]
        area_mean[a] = np.nanmean(np.array(rows), axis=0)
    ar = procedure({u: sim["margin"][u] - area_mean[st[u].area] for u, b, d in ex3}, brand_change)
    ok(ar["call"] == "TAL" and bins(ar["shares"]) == gold_bins, "SF7 the area-relative margin against its own changes")
    sd_route = cross_section(ex3, np.mean, lambda x: x.std(ddof=1))
    ok(AN.inquiry(sd_route) is None and sum(f for f, n in sd_route.values()) == 0,
       "ladder: a cross-section z-score on the plain standard deviation, which counts the Tallis stations in its own "
       "scale, finds no station")
    coarse = procedure(sim["margin"], brand_change, sel=lambda ks: (ks - CH) % 28 == 0)
    nref = sum(1 for k in M.placebo_indices() if (k - CH) % 28 == 0)
    ok(coarse["call"] is None and 1.0 / (nref + 1) > W.FDR * 28 / len(ex3),
       "ladder: a placebo of one date in 28 cannot reach Benjamini-Hochberg's bar for %d stations" % len(ex3))
    tal = [u for u, b, d in ex3 if b == "TAL"]
    found_uids = {x[0] for x, f in zip(gold["ex"], gold["found"]) if f}
    ok(found_uids == {u for u in tal if st[u].brand_priced_on(W.CHANGE)}, "every station found is a station Tallis "
                                                                           "prices itself, and every such station is found")
    ok({st[u].area for u in found_uids} == set(M.TALLIS_OWN), "the found stations trade in Tallis's four towns")
    found_d = [d for u, b, d in ex3 if u in found_uids]
    others_d = [d for u, b, d in ex3 if u not in found_uids]
    ok(min(found_d) - max(others_d) >= 1.0, "the found stations' changes clear every other station's by a cent a litre")
    ok(2.5 <= np.mean(found_d) <= 3.5, "the found stations' margins rose about 3 cents a litre")
    lvl = [np.nanmean(sim["margin"][u][W.DAY_INDEX[dt.date(2026, 2, 1)]:]) -
           np.nanmean(sim["margin"][u][:W.DAY_INDEX[dt.date(2025, 12, 1)]]) for u in found_uids]
    ok(2.5 <= np.mean(lvl) <= 3.5, "and stayed there through 2026 (%.2f)" % np.mean(lvl))
    first_of = [d for d in [dt.date(2025, m, 1) for m in range(2, 12)] + [dt.date(2026, m, 1) for m in range(3, 11)]]
    med = collections.defaultdict(list)
    for d in first_of:
        k = W.DAY_INDEX[d]
        p = []
        for u, b, _ in ex3:
            m = sim["margin"][u]
            a, z = m[k:k + 28], m[k - 28:k]
            a, z = a[~np.isnan(a)], z[~np.isnan(z)]
            p.append(stats.ttest_ind(a, z, equal_var=False, alternative="greater")[1] if len(a) >= 20 and len(z) >= 20
                     else 1.0)
        sh = AN.shares(ex3, AN.bh(np.array(p)))
        for b in W.CODES:
            med[b].append(rate(sh, b))
    ok(all(np.median(med[b]) >= 0.2 for b in W.CODES), "the easy test at the first of every month away from the change "
                                                         "finds a fifth or more of every brand's stations (medians %s)"
       % {b: round(float(np.median(v)), 2) for b, v in med.items()})
    for a in ("BRK", "POS"):
        c = sim["cyc"][a]
        jumps = [W.DAYS[k] for k in range(1, M.N) if c[k] - c[k - 1] > 1.0]
        ok(any(dt.date(2025, 12, 26) <= d <= dt.date(2026, 1, 5) for d in jumps),
           "%s restored its prices around the turn of the year" % a)
    hbl_areas = collections.Counter(st[u].area for u, b, d in ex3 if b == "HBL")
    ok(hbl_areas["BRK"] + hbl_areas["POS"] > 0.8 * sum(hbl_areas.values()), "most Harbourline stations trade in "
                                                                              "Brackwater and Port Ostley")

    # ---- semantic forks
    ok("06:00 to 22:00 local time" in standard and "no price in force at 06:00" in standard, "SF1 the daily price")
    for kind in ("noon", "mean_posted", "tw24"):
        mv = {s.uid: AN.margins(daily_variant(sim["events"][s.uid], s.closed, kind), rk) for s in sim["stations"]}
        same_call(procedure(mv, brand_change), "SF1 the daily price as %s" % kind)
    ok("net of the 5 per cent sales tax" in standard, "SF2 the margin net of the sales tax")
    tax = {u: sim["daily"][u] / 10.0 - AN.LEVY - rk for u in sim["margin"]}
    same_call(procedure(tax, brand_change), "SF2 the margin with tax")
    tax95 = {u: sim["daily"][u] / 10.0 * 0.95 - AN.LEVY - rk for u in sim["margin"]}
    same_call(procedure(tax95, brand_change), "SF2 the tax taken off as 5 per cent of the price")
    ok("the 28 days from the change less the mean of its daily margins over the 28 days before it" in standard,
       "SF3 the windows")
    ok("at least 20 of each of the two 28-day periods" in standard, "SF4 the examined stations")
    allb = collections.Counter(s.brand_on(W.CHANGE) for s in sim["stations"] if s.open_on(W.CHANGE))
    for b in W.CODES:
        ok(DL.tens((shares[b][0], allb[b])) == gold_bins[b], "SF4 and SF9 the share over every station carrying %s "
                                                              "on the day of the change" % b)
    ok(bins(variants["placebo, two-sided"]["shares"]) == gold_bins, "SF5 a two-sided test files the same shares")
    ok("by the Benjamini-Hochberg procedure" in standard, "SF6 the standard names Benjamini-Hochberg")
    ok(bins(mc["valid_procs"]["storey"]) == gold_bins, "SF6 Storey's adaptive procedure files the same shares")
    ok("on the station register on the day of the change" in standard, "SF8 the brand on the day of the change")
    ok("largest share of its examined stations" in standard, "SF9 the share over examined stations")
    for b in W.CODES:
        x = 100.0 * shares[b][0] / shares[b][1]
        ok(abs((x % 10) - 5) >= 2.0 or x == 0, "SF10 %s's share clear of a rounding edge (%.1f)" % (b, x))
    top = sorted((rate(shares, b) for b in W.CODES), reverse=True)
    ok(top[0] >= 3 * W.INQUIRY_SHARE - 0.01 and top[1] <= 0.05, "SF11 one brand at three times the bar, none other "
                                                                 "above one in twenty")
    ok("its latest published price applies" in standard, "SF12 the rack carried")
    ok("UTC-5, and UTC-4" in notes, "SF13 the clock")

    # ---- each load-bearing fact in exactly one file
    for frag, where in (("Benjamini-Hochberg", D.STANDARD), ("06:00 to 22:00", D.STANDARD), ("25.38", D.SCHEDULE),
                        ("UTC-4", D.NOTES), ("one in ten", D.STANDARD)):
        ok([f for f, t in texts.items() if frag in t] == [where], "%r stated once, in %s" % (frag, where))

    # ---- the asks
    rets, fr = sim["returns"], sim["frame"]
    lat = latest_returns(rets)
    holder = {}
    for s in sim["stations"]:
        if not getattr(s, "reissued_to", None):
            holder[s.sfn] = s
    gV = {W.NAME[b]: v for b, v in {b: [SV.thousands(x) for x in sim["litres"][b]] for b in W.CODES}.items()}
    gold_v = {b: gV[W.NAME[b]] for b in W.CODES}
    ok(gold_v == {b: [fig["regular_litres_thousands"][W.NAME[b]][m] for m in DL.MONTHS] for b in W.CODES},
       "V figures.json carries the litres")
    truth = collections.defaultdict(int)
    for r in lat:
        s = st[r["uid"]]
        if r["ref"] in SV.ASK_MONTHS and s.brand_on(max(r["ref"], s.opened)) in W.CODES:
            truth[(s.brand_on(max(r["ref"], s.opened)), r["ref"])] += r["reg"]
    ok({b: [SV.thousands(truth[(b, m)]) for m in SV.ASK_MONTHS] for b in W.CODES} == gold_v, "V the golden is the truth")
    for b in W.CODES:
        for x in sim["litres"][b]:
            ok(x % 1000 != 500, "V no litre total on a half thousand")
    today = lambda s, m: s.brands[-1][0]
    joined = lambda r: holder[r["sfn"]].frame_added <= SV.month_end(r["ref"])
    stops = {
        "the cracker: latest revision, reference month, the frame's holder, its brand that month":
            litres_stop(sim, lat, "ref", lambda r: holder[r["sfn"]], lambda s, m: s.brand_on(max(m, s.opened))),
        "the over-cleaner: returns dated before the holder joined the frame dropped":
            litres_stop(sim, [r for r in lat if joined(r)], "ref", lambda r: holder[r["sfn"]],
                        lambda s, m: s.brand_on(max(m, s.opened))),
        "latest revision, reference month, the frame's holder today":
            litres_stop(sim, lat, "ref", lambda r: holder[r["sfn"]], today),
        "every row, processing month, the frame's holder today":
            litres_stop(sim, rets, "proc", lambda r: holder[r["sfn"]], today),
    }
    for k, v in stops.items():
        hits = [(b, i) for b in W.CODES for i in range(3) if v[b][i] == gold_v[b][i]]
        ok(not hits, "V %s misses every cell (%s)" % (k, hits))
    in_month = lambda r, m: st[r["uid"]].brand_on(max(m, st[r["uid"]].opened))
    cell = lambda b, m: (b, SV.ASK_MONTHS.index(m))
    rev_cells = {cell(in_month(r, r["ref"]), r["ref"]) for r in rets if r["rev"] > 1 and r["ref"] in SV.ASK_MONTHS
                 and in_month(r, r["ref"]) in W.CODES}
    late_cells = set()
    for r in lat:
        if r["proc"] != r["ref"]:
            for m in (r["ref"], r["proc"]):
                if m in SV.ASK_MONTHS and in_month(r, r["ref"]) in W.CODES:
                    late_cells.add(cell(in_month(r, r["ref"]), m))
    hazards = {
        "every revision summed, reference month, the right station":
            (litres_stop(sim, rets, "ref", lambda r: st[r["uid"]], lambda s, m: s.brand_on(max(m, s.opened))),
             rev_cells),
        "latest revision, processing month, the right station":
            (litres_stop(sim, lat, "proc", lambda r: st[r["uid"]], lambda s, m: s.brand_on(max(m, s.opened))),
             late_cells),
    }
    for k, (v, cells) in hazards.items():
        moved = {(b, i) for b in W.CODES for i in range(3) if v[b][i] != gold_v[b][i]}
        ok(moved == cells and len(cells) >= 12, "V hazard %s moves exactly the %d cells it touches" % (k, len(cells)))
    hz = litres_stop(sim, lat, "ref", lambda r: st[r["uid"]], today)
    touched = {(b, i) for b in W.CODES for i in range(3) if hz[b][i] != gold_v[b][i]}
    reb = set()
    for s in sim["stations"]:
        if len(s.brands) > 1 and s.brands[1][1] > SV.ASK_MONTHS[0]:
            for i, m in enumerate(SV.ASK_MONTHS):
                if s.open_on(m) and s.brand_on(m) != s.brands[-1][0]:
                    reb.add((s.brand_on(m), i))
                    reb.add((s.brands[-1][0], i))
    ok(touched == {x for x in reb if x[0] in W.CODES}, "V the rebranding hazard moves exactly the cells it touches")
    late = sum(1 for r in rets if r["proc"] != r["ref"])
    revs = sum(1 for r in rets if r["rev"] > 1)
    ok(late >= 100 and revs >= 150, "V hazards: %d returns processed in another month, %d revisions" % (late, revs))
    old_holders = [s for s in sim["stations"] if getattr(s, "reissued_to", None)]
    ok(len(old_holders) == fig["survey_references_reissued"] == 8, "V eight references reissued")
    reg = openpyxl.load_workbook(os.path.join(target, R.REGISTER), read_only=True)
    sites = collections.Counter(V.norm(r[2] + " " + r[3]) for r in list(reg["Stations"].iter_rows(values_only=True))[1:])
    ok(max(sites.values()) == 1, "V every register site address is unique")
    ok(len({s.closed for s in old_holders}) == len(old_holders), "V the closing dealers close on different days")
    for s in old_holders:
        n = st[s.reissued_to]
        ok(dt.date(2026, 3, 2) <= s.closed < dt.date(2026, 8, 1) and n.frame_added >= dt.date(2026, 8, 1) and
           n.opened <= W.LOG_START and n.open_on(SV.ASK_MONTHS[0]),
           "V the reference was freed at a closure and given to a station trading all year that joined the frame "
           "from August")
        ok(not any(r["uid"] == n.uid and r["ref"] < n.frame_added.replace(day=1) for r in rets),
           "V the new holder filed nothing before it joined the frame")
        ok(sites[V.norm(SV.site_address(s))] == 1, "V the closed station's site address matches one register station")
        for m in SV.ASK_MONTHS:
            ok(not any(W.NAME[b].casefold() in SV.reporting_name(s, m).casefold() for b in W.CODES),
               "V the closed dealer's returns do not name its brand")
        ok(s.brands[0][0] != st[s.reissued_to].brands[0][0] or s.brands[0][0] == W.INDEPENDENT,
           "V each reissued reference changes brand")
    ok(sum(1 for s in old_holders if s.brands[0][0] == W.INDEPENDENT) == 1, "V one inert instance (unbranded)")
    # P: the stations each brand prices itself on 1 December 2026
    gp = {b: fig["priced_by_brand_2026_12_01"][W.NAME[b]] for b in W.CODES}
    truth_p = {b: sum(1 for s in sim["stations"] if s.open_on(W.SCOPE_DATE) and s.brand_on(W.SCOPE_DATE) == b
                      and s.brand_priced_on(W.SCOPE_DATE)) for b in W.CODES}
    ok(gp == truth_p, "P the evidence reproduces the truth exactly")
    own = {b: sum(1 for s in sim["stations"] if s.open_on(W.SCOPE_DATE) and s.brand_on(W.SCOPE_DATE) == b
                  and s.ownership == "company") for b in W.CODES}
    ok(all(own[b] != gp[b] for b in W.CODES), "P the cracker (the register's site ownership) misses every brand %s" % own)
    api = priced_variant(sim, sim["subs"], need=1, share=0.01)
    ok(all(api[b] != gp[b] for b in W.CODES), "P every station with an API submission misses every brand %s" % api)
    own_api = {b: sum(1 for s in sim["stations"] if s.open_on(W.SCOPE_DATE) and s.brand_on(W.SCOPE_DATE) == b and
                      s.ownership == "company" and any(r["uid"] == s.uid and r["channel"] == "API" for r in sim["subs"]))
               for b in W.CODES}
    ok(sum(own_api[b] != gp[b] for b in W.CODES) >= 5, "P company-owned with API submissions misses five or more "
                                                       "brands %s" % own_api)
    for kw in ({"need": 2}, {"need": 4}, {"share": 0.8}, {"share": 1.0}, {"share": 0.3},
               {"grades": ("RUL", "PUL", "ULSD")}, {"channel": None, "need": 4}, {"by": "minute", "need": 2},
               {"by": "minute", "need": 3}, {"by": "minute", "need": 2, "channel": None, "share": 0.4}):
        ok(priced_variant(sim, sim["subs"], **kw) == gp, "P the evidence read with %s converges" % kw)
    conv = [s for s in sim["stations"] if len(s.models) > 1]
    ok(len(conv) == 2 and all(s.models[1][1] < dt.date(2026, 11, 1) for s in conv), "P two pricing conversions before "
                                                                                     "November")
    # separation: no ask row in the main population
    ok(all(s.closed is None or s.closed >= W.CHANGE + dt.timedelta(days=28) for s in old_holders),
       "separation: the reissued references' closed stations trade through the change windows")
    ok(all(s.frame_added > W.CHANGE + dt.timedelta(days=28) for s in sim["stations"] if getattr(s, "reissue_of", None)),
       "separation: the reissued references' new holders join the frame after the change windows")
    with open(os.path.join(target, R.COMPLAINTS), encoding="utf-8") as fh:
        head = fh.readline()
    ok("brand" not in head, "the complaints register names no brand")

    # ---- the literal read of the examined population: the register's brand on the day and the price standing until
    # the next change or the closing day
    hist = list(reg["Brand history"].iter_rows(values_only=True))[1:]
    rows_st = {r[0]: r for r in list(reg["Stations"].iter_rows(values_only=True))[1:]}
    for sid, b, f, to in hist:
        cl = rows_st[sid][6]
        ok(cl is None or (to is not None and to < cl), "the register ends every closed station's brand before it closed")
    lit = {sid for sid, b, f, to in hist if b in W.CODES and f.date() <= W.CHANGE and (to is None or W.CHANGE <= to.date())}
    num_of = {s.uid: R.register_number(s) for s in sim["stations"]}
    ex_lit = {u for u in st if num_of[u] in lit and not np.isnan(AN.window_change(sim["margin"][u], CH))}
    ok(ex_lit == {u for u, b, d in ex3} and len(ex_lit) == fig["stations_examined"],
       "the literal read examines the same %d stations" % fig["stations_examined"])
    ok("or until\n  the station closes" in notes or "or until the station closes" in notes, "the notes end a price at closing")

    # ---- dates and the notes
    ok({f for f in files if f in notes} == set(files) - {D.NOTES}, "the extract notes name every shipped file")
    lateness = []
    for f in files:
        if f.endswith((".csv", ".md", ".json")):
            t = open(os.path.join(target, f), encoding="utf-8").read()
            if any(x > W.AS_OF.isoformat() for x in set(re.findall(r"\b20\d\d-\d\d-\d\d\b", t))):
                lateness.append(f)
    ok(lateness == [], "no record runs past the extract date (%s)" % lateness)

    # ---- realism sweep
    ex_counts = [shares[b][1] for b in W.CODES]
    ok(len(set(ex_counts)) >= 5, "examined counts vary by brand")
    ok(len({x for b in W.CODES for x in gold_v[b]}) == 21, "every litre figure distinct")
    ok(len(set(gp.values())) >= 5, "priced-by-brand counts vary")
    ok(len({r["addr"] for r in rets}) > 380 and sum(1 for s in sim["stations"] if s.style == 1) > 50,
       "site addresses vary in style")
    ops = collections.Counter(SV.operator(s) for s in sim["stations"] if s.model not in ("COCO",))
    ok(max(ops.values()) <= 5, "no operator name runs more than five stations")
    pr = [float(r["price_cents_per_litre"]) for r in log[::50]]
    ok(85 < min(pr) and max(pr) < 150, "pump prices in a province's range")
    ok(len({r["station_number"] for r in log}) > 400, "four hundred stations on the log")
    for k, (who, role) in W.PEOPLE.items():
        if k != "commissioner":
            ok(any(who in t for t in texts.values()), "persona %s in the pack" % who)

    # ---- goldens agree with figures.json
    ok(sorted(os.listdir(golden)) == sorted([DL.MEMO, DL.SCOPE]), "goldens are the two files asked")
    from docx import Document
    doc = Document(os.path.join(golden, DL.MEMO))
    paras = [p.text for p in doc.paragraphs]
    body = " ".join(paras)
    ok("Tallis Petroleum for the 2027 pass-through inquiry" in paras, "memo title is the call")
    ok("examine Tallis Petroleum" in body, "memo names the brand")
    t = doc.tables[1]
    rows = [[c.text for c in r.cells] for r in t.rows]
    ok(rows[0] == ["Brand", "Stations examined", "Share found"], "memo table header")
    for r, b in zip(rows[1:], DL.ORDER):
        nm = W.NAME[b]
        ok(r == [nm, "%d" % shares[b][1], "%s per cent" % fig["share_found_nearest_ten_per_cent"][nm]], "memo row " + nm)
    ok("%d for Tallis on 1 December 2026" % gp["TAL"] in body, "memo carries Tallis's priced-by-brand count")
    ok("all %d stations examined" % fig["stations_examined"] in body, "memo carries the examined total")
    import zipfile
    with zipfile.ZipFile(os.path.join(golden, DL.MEMO)) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
        core = z.read("docProps/core.xml").decode()
    ok(len(media) == 1, "one chart in the memo")
    ok("Kendra Andrews" in core and "python-docx" not in core, "memo container scrubbed")
    for p in paras:
        ok("\u2014" not in p and "\u2013" not in p, "no dash in the memo")
    wb = openpyxl.load_workbook(os.path.join(golden, DL.SCOPE), read_only=True)
    ok(wb.sheetnames == ["Inquiry scope", "Notes"], "workbook sheets named")
    got = {r[1]: r for r in wb["Inquiry scope"].iter_rows(min_row=5, values_only=True) if r and r[1] in W.CODES}
    for b in W.CODES:
        ok(list(got[b][2:6]) == [gp[b]] + gold_v[b], "workbook row " + b)
    for r in wb["Notes"].iter_rows(values_only=True):
        for c in r:
            ok(c is None or ("\u2014" not in c and "\u2013" not in c), "no dash in the workbook notes")
    ok(min(VALID.values()) >= 26 and max(VALID.values()) <= 32, "every valid null finds 26 to 32 Tallis stations %s"
       % VALID)
    others = {k: v for k, v in VALID.items()}
    print("valid nulls, Tallis stations found: %s" % others)
    print("checks: %d assertions passed (call %s; shares %s; default path %s; examined %d, found %d)" % (
        N[0], W.NAME[mc["call"]], {b: gold_bins[b] for b in W.CODES}, {b: DL.tens(w[b]) for b in W.CODES},
        fig["stations_examined"], fig["stations_found"]))
