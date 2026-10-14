"""The pack's data files, written from the simulated survey: both collection systems' exports, the December 2026
release, the fourth-quarter fieldwork log and its unit lookup, the region concordance, the weighting benchmarks,
the seasonal factors, the response report, the file dictionary and the export manifest."""
import csv
import datetime as dt

import numpy as np
import pyarrow as pa
import pyarrow.parquet as pq

import world as W

FILES = {
    "old": "lfs_cai2014_estimation_cells_2022-01_2025-11.csv",
    "new": "lfs_ics_estimation_cells_2025-07_2026-12.csv",
    "release": "lfs_release_2026-12_tables.xlsx",
    "field": "lfs_fieldwork_contacts_2026q4.parquet",
    "psu": "lfs_psu_region_lookup_2025.csv",
    "conc": "region_classification_concordance_2019_2025.csv",
    "pop": "population_benchmarks_15plus_2025-2027.csv",
    "seas": "lfs_seasonal_factors_2027_projected.csv",
    "resp": "lfs_response_rates_2025-2026.xlsx",
    "dict": "lfs_file_dictionary.csv",
    "manifest": "export_manifest.csv",
    "note": "lfs_redesigned_questionnaire_methodology_note.pdf",
    "act": "regional_employment_support_act_2021_extract.pdf",
}
AGE_CODE = {a: i + 1 for i, a in enumerate(W.AGES)}


def _csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(header)
        n = 0
        for r in rows:
            w.writerow(r)
            n += 1
    return n


def region_code_in(t, code25):
    """the region code an export carries in month t: the 2019 classification to December 2024, the 2025 one after"""
    return W.CODE19[code25] if W.MONTHS[t] < "2025-01" else code25


def write_old(S, path):
    """the previous questionnaire's estimation cells from CAI 2014: long, one row per labour force status, weighted
    persons in thousands to three decimals, months as YYYYMM"""
    rows = []
    for t in range(W.T):
        if W.MONTHS[t] > "2025-11":
            continue
        for r, code in enumerate(W.RCODE):
            for g in range(6):
                if W.redesigned(t, g):
                    continue
                for s in range(2):
                    for a in range(len(W.AGES)):
                        emp, une, nilf = S["emp"][t, r, g, s, a], S["une"][t, r, g, s, a], S["nilf"][t, r, g, s, a]
                        lf, u, n = S["lf"][t, r, g, s, a], S["u"][t, r, g, s, a], S["n"][t, r, g, s, a]
                        for stat, resp, wtd in ((1, lf - u, emp), (2, u, une), (3, n - lf, nilf)):
                            rows.append((W.MONTHS[t].replace("-", ""), region_code_in(t, code), W.GROUPS[g], s + 1,
                                         a + 1, stat, int(resp), "%.3f" % (wtd / 1000.0)))
    return _csv(path, ["REF_MTH", "REGION", "ROTGRP", "SEX", "AGEGRP", "LFSTAT", "RESP", "WTD000"], rows)


def write_new(S, path):
    """the redesigned questionnaire's estimation cells from ICS: wide, weighted persons"""
    rows = []
    for t in range(W.PHASE, W.T):
        for r, code in enumerate(W.RCODE):
            for g in range(6):
                if not W.redesigned(t, g):
                    continue
                for s in range(2):
                    for a in range(len(W.AGES)):
                        rows.append((W.MONTHS[t], code, W.GROUPS[g], W.SEX[s], W.AGES[a], int(S["n"][t, r, g, s, a]),
                                     int(S["emp"][t, r, g, s, a]), int(S["une"][t, r, g, s, a]),
                                     int(S["nilf"][t, r, g, s, a])))
    return _csv(path, ["period", "region", "rotation_group", "sex", "age_band", "respondents", "employed",
                       "unemployed", "not_in_labour_force"], rows)


def published_tables(S):
    """the release's figures, computed from the cells both systems export, exactly as the release prints them"""
    une = S["une"].sum(axis=(2, 3, 4)).astype(float) / 6.0
    emp = S["emp"].sum(axis=(2, 3, 4)).astype(float) / 6.0
    lab = une + emp
    reg = np.round(100.0 * une / lab, 1)
    nat_rate = 100.0 * une.sum(1) / lab.sum(1)
    sa_f = seasonal_factors_hist()
    return dict(reg=reg, reg_lf=np.round(lab / 1000.0, 1), nat_lf=np.round(lab.sum(1) / 1000.0, 1),
                nat_emp=np.round(emp.sum(1) / 1000.0, 1), nat_une=np.round(une.sum(1) / 1000.0, 1),
                nat=np.round(nat_rate, 1), nat_sa=np.round(nat_rate - sa_f, 1), sa_f=sa_f)


def seasonal_factors_hist():
    rng = np.random.default_rng(W.SEED + 5)
    f = np.array([W.SEAS[t % 12] for t in range(W.T)]) + rng.normal(0, 0.02, W.T)
    return np.round(f, 2)


def write_release(P, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Contents"
    bold = Font(name="Arial", bold=True, size=11)
    norm = Font(name="Arial", size=9)
    lines = [(W.SURVEY + ", December 2026", bold), ("%s. Released 21 January 2027." % W.OFFICE, norm), ("", norm),
             ("Table 1  Labour force status, persons aged 15 and over, monthly, original and seasonally adjusted", norm),
             ("Table 2  Unemployment rate by region, monthly, original", norm),
             ("Table 3  Labour force by region, monthly, original", norm), ("", norm),
             ("Estimates are averages of the six rotation groups' estimates. Regions are on the 2025 regional "
              "classification throughout.", norm),
             ("July to November 2025 combine rotation groups answering the previous questionnaire and rotation groups "
              "answering the redesigned questionnaire; from December 2025 every rotation group answers the redesigned "
              "questionnaire. Earlier months have not been revised.", norm),
             ("Contact: Labour Market Statistics, %s." % W.OFFICE, norm)]
    for i, (txt, f) in enumerate(lines, 1):
        if txt:
            ws.cell(row=i, column=1, value=txt).font = f
    ws.column_dimensions["A"].width = 118

    def sheet(title, heading, header, rows, fmt, widths, notes):
        s = wb.create_sheet(title)
        s.cell(row=1, column=1, value=heading).font = bold
        s.cell(row=2, column=1, value=W.SURVEY + ", December 2026").font = norm
        for j, h in enumerate(header, 1):
            c = s.cell(row=4, column=j, value=h)
            c.font = Font(name="Arial", bold=True, size=9)
            c.alignment = Alignment(wrap_text=True, vertical="bottom")
        for i, r in enumerate(rows, 5):
            for j, v in enumerate(r, 1):
                c = s.cell(row=i, column=j, value=v)
                c.font = norm
                if j > 1:
                    c.number_format = fmt if isinstance(fmt, str) else fmt[j - 2]
        k = 5 + len(rows) + 1
        for n in notes:
            s.cell(row=k, column=1, value=n).font = Font(name="Arial", size=8, italic=True)
            k += 1
        for j, wd in enumerate(widths, 1):
            s.column_dimensions[chr(64 + j)].width = wd
        s.freeze_panes = "B5"
        return s

    mlabel = lambda t: dt.date(int(W.MONTHS[t][:4]), int(W.MONTHS[t][5:]), 1).strftime("%b-%Y")
    r1 = [[mlabel(t), float(P["nat_lf"][t]), float(P["nat_emp"][t]), float(P["nat_une"][t]), float(P["nat"][t]),
           float(P["nat_sa"][t])] for t in range(W.T)]
    sheet("Table 1", "Table 1  Labour force status, persons aged 15 and over",
          ["Month", "Labour force ('000)", "Employed ('000)", "Unemployed ('000)", "Unemployment rate, original (%)",
           "Unemployment rate, seasonally adjusted (%)"], r1, ["#,##0.0", "#,##0.0", "#,##0.0", "0.0", "0.0"],
          [11, 13, 13, 13, 15, 17],
          ["July to November 2025: both questionnaires in the field (see Contents).",
           "Seasonally adjusted series are re-estimated each month; the original series is not revised."])
    r2 = [[mlabel(t)] + [float(P["reg"][t, r]) for r in range(W.R)] for t in range(W.T)]
    sheet("Table 2", "Table 2  Unemployment rate by region, original (%)",
          ["Month"] + ["%s %s" % (c, W.RNAME[c]) for c in W.RCODE], r2, "0.0", [11] + [12] * W.R,
          ["Regions on the 2025 regional classification.", "July to November 2025: both questionnaires in the field."])
    r3 = [[mlabel(t)] + [float(P["reg_lf"][t, r]) for r in range(W.R)] for t in range(W.T)]
    sheet("Table 3", "Table 3  Labour force by region, original ('000)",
          ["Month"] + ["%s %s" % (c, W.RNAME[c]) for c in W.RCODE], r3, "#,##0.0", [11] + [12] * W.R,
          ["Regions on the 2025 regional classification."])
    wb.save(path)
    return 3 * W.T


def write_fieldwork(rows, path):
    """the fourth quarter of 2026 as the fieldwork system logs it: one row per sample slot and month"""
    cols = {"contact_month": [], "psu": [], "rotation_group": [], "line": [], "dwelling_ref": [], "interview_no": [],
            "outcome": [], "mode": [], "interviewer_area": []}
    order = sorted(rows, key=lambda r: (r[0], r[8], r[1], r[2], r[3]))
    for r in order:
        cols["contact_month"].append(r[0])
        cols["psu"].append(r[1])
        cols["rotation_group"].append(r[2])
        cols["line"].append(r[3])
        cols["dwelling_ref"].append("%s%s%02d" % (r[1], r[2], r[3]))
        cols["interview_no"].append(r[4])
        cols["outcome"].append(r[6])
        cols["mode"].append(r[7])
        cols["interviewer_area"].append(r[8])
    schema = pa.schema([("contact_month", pa.string()), ("psu", pa.string()), ("rotation_group", pa.string()),
                        ("line", pa.int8()), ("dwelling_ref", pa.string()), ("interview_no", pa.int8()),
                        ("outcome", pa.int16()), ("mode", pa.string()), ("interviewer_area", pa.string())])
    tab = pa.table(cols, schema=schema)
    pq.write_table(tab, path, compression="snappy", use_dictionary=True, write_statistics=False, version="2.6",
                   data_page_version="1.0")
    return len(order)


def write_psu(numbers, reg, path):
    rows = sorted(("%04d" % p, W.RCODE[r], W.RNAME[W.RCODE[r]]) for p, r in zip(numbers, reg))
    return _csv(path, ["psu", "region_2025", "region_name"], rows)


def write_concordance(path):
    rows = sorted((W.CODE19[c], W.RNAME[c], c, W.RNAME[c]) for c in W.RCODE)
    return _csv(path, ["region_2019", "region_2019_name", "region_2025", "region_2025_name"], rows)


def write_population(pop, path):
    """the population benchmarks the rotation groups are weighted to, January of each year, with the 2027 projection"""
    rows = []
    for y in (2025, 2026, 2027):
        k = min(W.MONTHS.index("%d-01" % min(y, 2026)), W.T - 1)
        grow = 1.0 + (0.008 if y == 2027 else 0.0)
        for r, c in enumerate(W.RCODE):
            for s in range(2):
                for a in range(len(W.AGES)):
                    rows.append((y, c, W.SEX[s], W.AGES[a], int(round(pop[k, r, s, a] * grow, -1))))
    return _csv(path, ["year", "region_2025", "sex", "age_band", "persons_1_january"], rows)


def write_seasonal(path):
    rng = np.random.default_rng(W.SEED + 7)
    rows = [("2027-%02d" % m, "%.2f" % (W.SEAS[m - 1] + rng.normal(0, 0.03))) for m in range(1, 13)]
    return _csv(path, ["month", "projected_factor_points"], rows)


def write_response(S, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Response"
    ws.cell(row=1, column=1, value="%s: response by region, January 2025 to December 2026" % W.SURVEY).font = Font(
        name="Arial", bold=True, size=11)
    ws.cell(row=2, column=1, value="Responding dwellings as a percentage of eligible dwellings (vacant, derelict and "
                                   "non-private dwellings excluded), all rotation groups. Survey Operations.").font = Font(
        name="Arial", size=9)
    hdr = ["Month"] + [W.RNAME[c] for c in W.RCODE] + ["All regions"]
    for j, h in enumerate(hdr, 1):
        ws.cell(row=4, column=j, value=h).font = Font(name="Arial", bold=True, size=9)
    k = 5
    for t in range(W.MONTHS.index("2025-01"), W.T):
        d, e = S["D"][t].sum(1), S["E"][t].sum(1)
        vals = 100.0 * d / e
        ws.cell(row=k, column=1, value=W.MONTHS[t])
        for j, v in enumerate(vals, 2):
            c = ws.cell(row=k, column=j, value=round(float(v), 1))
            c.number_format = "0.0"
        allv = 100.0 * d.sum() / e.sum()
        ws.cell(row=k, column=W.R + 2, value=round(float(allv), 1)).number_format = "0.0"
        k += 1
    rows = k - 5
    for txt in ("Eligible dwellings exclude dwellings found vacant (31), derelict or demolished (32) or not private "
                "dwellings (33). All six rotation groups.",
                "Response falls over the summer holidays in December and January, most in the Capital and Harbour."):
        k += 1
        ws.cell(row=k, column=1, value=txt).font = Font(name="Arial", size=8, italic=True)
    ws.freeze_panes = "B5"
    ws.column_dimensions["A"].width = 10
    wb.save(path)
    return rows


DICTIONARY = [
    ("lfs_cai2014_estimation_cells_2022-01_2025-11.csv", "", "Estimation cells from CAI 2014: one row per reference "
     "month, region, rotation group, sex, age group and labour force status, for the rotation groups that answered the "
     "previous questionnaire."),
    ("", "REF_MTH", "Reference month, YYYYMM."),
    ("", "REGION", "Region code on the regional classification in force in the reference month: the 2019 classification "
     "to December 2024, the 2025 classification from January 2025."),
    ("", "ROTGRP", "Rotation group, A to F."),
    ("", "SEX", "1 male, 2 female."),
    ("", "AGEGRP", "1 aged 15 to 24, 2 aged 25 to 34, 3 aged 35 to 44, 4 aged 45 to 54, 5 aged 55 and over."),
    ("", "LFSTAT", "1 employed, 2 unemployed, 3 not in the labour force."),
    ("", "RESP", "Responding persons in the cell with that status (unweighted)."),
    ("", "WTD000", "Weighted persons, thousands. Each rotation group is weighted to the population aged 15 and over."),
    ("lfs_ics_estimation_cells_2025-07_2026-12.csv", "", "Estimation cells from ICS: one row per period, region, "
     "rotation group, sex and age band, for the rotation groups that answered the redesigned questionnaire."),
    ("", "period", "Reference month, YYYY-MM."),
    ("", "region", "Region code, 2025 regional classification."),
    ("", "rotation_group", "Rotation group, A to F."),
    ("", "sex", "M or F."),
    ("", "age_band", "Age band in years."),
    ("", "respondents", "Responding persons aged 15 and over in the cell (unweighted)."),
    ("", "employed", "Weighted persons employed. Each rotation group is weighted to the population aged 15 and over."),
    ("", "unemployed", "Weighted persons unemployed."),
    ("", "not_in_labour_force", "Weighted persons not in the labour force."),
    ("lfs_fieldwork_contacts_2026q4.parquet", "", "Fieldwork outcomes, October to December 2026: one row per sample "
     "reference and contact month."),
    ("", "contact_month", "Survey month of the contact, YYYY-MM."),
    ("", "psu", "Primary sampling unit, four digits."),
    ("", "rotation_group", "Rotation group, A to F."),
    ("", "line", "Line within the primary sampling unit and rotation group, numbered from 1."),
    ("", "dwelling_ref", "Sample reference: primary sampling unit, rotation group and line."),
    ("", "interview_no", "Interview number, 1 to 6 (1 is the first interview)."),
    ("", "outcome", "11 full interview, 12 partial interview, 21 refusal, 22 no contact, 23 other non-interview, "
     "31 vacant, 32 derelict or demolished, 33 not a private dwelling. Codes 11 and 12 are responding dwellings."),
    ("", "mode", "F2F face to face, TEL telephone."),
    ("", "interviewer_area", "Interviewer area responsible for the contact."),
    ("lfs_psu_region_lookup_2025.csv", "", "Primary sampling units and their region, 2025 regional classification."),
    ("region_classification_concordance_2019_2025.csv", "", "Region codes on the 2019 and 2025 regional "
     "classifications. The 2025 review renumbered the regions; their boundaries did not change."),
    ("population_benchmarks_15plus_2025-2027.csv", "", "Population aged 15 and over at 1 January by region, sex and "
     "age band, the benchmarks the rotation groups are weighted to; 2027 projected."),
    ("lfs_seasonal_factors_2027_projected.csv", "", "Projected additive seasonal factors for the national unemployment "
     "rate, 2027."),
    ("lfs_response_rates_2025-2026.xlsx", "", "Monthly response rates by region, Survey Operations."),
    ("lfs_release_2026-12_tables.xlsx", "", "Labour Force Survey, December 2026: the published tables."),
]


def write_dictionary(path):
    return _csv(path, ["file", "field", "description"], DICTIONARY)


def write_manifest(path, counts):
    rows = [
        (FILES["old"], "Estimation cells, previous questionnaire (CAI 2014)", "January 2022 to November 2025",
         counts.get("old", ""), "Labour Market Statistics", W.EXTRACT),
        (FILES["new"], "Estimation cells, redesigned questionnaire (ICS)", "July 2025 to December 2026",
         counts.get("new", ""), "Labour Market Statistics", W.EXTRACT),
        (FILES["release"], "Labour Force Survey, December 2026, published tables", "January 2022 to December 2026",
         "", "Labour Market Statistics", "2027-01-21"),
        (FILES["field"], "Fieldwork outcomes", "October to December 2026", counts.get("field", ""),
         "Survey Operations", W.EXTRACT),
        (FILES["psu"], "Primary sampling units by region", "2025 regional classification", counts.get("psu", ""),
         "Survey Operations", W.EXTRACT),
        (FILES["conc"], "Regional classification concordance", "2019 and 2025 classifications",
         counts.get("conc", ""), "Classifications", "2024-11-29"),
        (FILES["pop"], "Population benchmarks used in weighting", "1 January 2025 to 2027", counts.get("pop", ""),
         "Population Statistics", "2026-11-12"),
        (FILES["seas"], "Projected seasonal factors, national unemployment rate", "2027", counts.get("seas", ""),
         "Time Series Analysis", "2027-01-08"),
        (FILES["resp"], "Response rates by region", "January 2025 to December 2026", "", "Survey Operations",
         W.EXTRACT),
        (FILES["dict"], "Field descriptions for the files above", "", "", "Labour Market Statistics", W.EXTRACT),
        (FILES["note"], "Methodology note on the redesigned questionnaire", "", "", "Labour Market Methods",
         "2025-06-16"),
        (FILES["act"], "Regional Employment Support Act 2021, extract", "consolidated to 1 July 2025", "",
         "Parliamentary Counsel", "2025-07-01"),
    ]
    return _csv(path, ["file", "content", "covers", "rows", "owner", "extracted"], rows)
