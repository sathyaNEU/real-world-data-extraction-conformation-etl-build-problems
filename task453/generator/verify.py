"""Independent verifier: reads only the shipped bundle and the goldens, on its own code path (csv, openpyxl and pyarrow
readers, a within-transformation fixed-effects fit), and recomputes the change, the restated trigger, the designation,
the regional series on the redesigned basis and the fourth-quarter dwelling count."""
import csv
import os
import re
import sys

import numpy as np

OLD = "lfs_cai2014_estimation_cells_2022-01_2025-11.csv"
NEW = "lfs_ics_estimation_cells_2025-07_2026-12.csv"
REL = "lfs_release_2026-12_tables.xlsx"
CONC = "region_classification_concordance_2019_2025.csv"
FIELD = "lfs_fieldwork_contacts_2026q4.parquet"
PSU = "lfs_psu_region_lookup_2025.csv"
PDF = "trigger_restatement_determination_2027.pdf"
CSVG = "regional_unemployment_redesigned_basis_2024-2026.csv"


def month_index(ym):
    y, m = int(ym[:4]), int(ym[5:7])
    return (y - 2022) * 12 + (m - 1)


def months_in_sample(ym, letter):
    """group A's dwellings are replaced in January and July, B's in February and August, ... (the note, section 1)"""
    m = int(ym[5:7])
    return ((m - 1 - "ABCDEF".index(letter)) % 6) + 1


def read_cells(target):
    to25 = {}
    with open(os.path.join(target, CONC), newline="") as f:
        for r in csv.DictReader(f):
            to25[r["region_2019"]] = r["region_2025"]
    cells = {}
    with open(os.path.join(target, OLD), newline="") as f:
        for r in csv.DictReader(f):
            ym = r["REF_MTH"][:4] + "-" + r["REF_MTH"][4:]
            reg = to25[r["REGION"]] if ym < "2025-01" else r["REGION"]
            k = (ym, reg, r["ROTGRP"], 0)
            u, l = cells.get(k, (0.0, 0.0))
            v = float(r["WTD000"]) * 1000.0
            if r["LFSTAT"] == "2":
                cells[k] = (u + v, l + v)
            elif r["LFSTAT"] == "1":
                cells[k] = (u, l + v)
            else:
                cells.setdefault(k, (u, l))
    with open(os.path.join(target, NEW), newline="") as f:
        for r in csv.DictReader(f):
            k = (r["period"], r["region"], r["rotation_group"], 1)
            u, l = cells.get(k, (0.0, 0.0))
            cells[k] = (u + float(r["unemployed"]), l + float(r["unemployed"]) + float(r["employed"]))
    return cells


def fit_effects(cells):
    """rate = month-region effect + month-in-sample effect (2..6) + the redesign's effect at each month in sample,
    weighted by labour force; fixed effects swept out by weighted demeaning within each month and region"""
    keys = sorted(cells)
    y = np.array([100.0 * cells[k][0] / cells[k][1] for k in keys])
    w = np.array([cells[k][1] for k in keys])
    mis = np.array([months_in_sample(k[0], k[2]) for k in keys])
    new = np.array([k[3] for k in keys])
    X = np.column_stack([(mis == j).astype(float) for j in range(2, 7)] +
                        [((mis == j) & (new == 1)).astype(float) for j in range(1, 7)])
    grp = {}
    for i, k in enumerate(keys):
        grp.setdefault((k[0], k[1]), []).append(i)
    yd, Xd = y.copy(), X.copy()
    for idx in grp.values():
        ww = w[idx] / w[idx].sum()
        yd[idx] -= (ww * y[idx]).sum()
        Xd[idx] -= (ww[:, None] * X[idx]).sum(0)
    sw = np.sqrt(w / w.mean())
    beta = np.linalg.lstsq(Xd * sw[:, None], yd * sw, rcond=None)[0]
    return beta[5:]


def read_release(target):
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(target, REL), read_only=True, data_only=True)
    head = [c for c in next(wb["Table 2"].iter_rows(min_row=4, max_row=4, values_only=True))]
    regs = [h.split(" ")[0] for h in head[1:] if h]
    rates = {}
    for row in wb["Table 2"].iter_rows(min_row=5, values_only=True):
        if row[0] and re.match(r"[A-Z][a-z]{2}-\d{4}$", str(row[0])):
            rates[str(row[0])] = [float(v) for v in row[1:1 + len(regs)]]
    nat = {}
    for row in wb["Table 1"].iter_rows(min_row=5, values_only=True):
        if row[0] and re.match(r"[A-Z][a-z]{2}-\d{4}$", str(row[0])):
            nat[str(row[0])] = float(row[4])
    return regs, rates, nat


def q4_dwellings(target):
    import pyarrow.parquet as pq
    t = pq.read_table(os.path.join(target, FIELD)).to_pydict()
    reg = {}
    with open(os.path.join(target, PSU), newline="") as f:
        for r in csv.DictReader(f):
            reg[r["psu"]] = r["region_2025"]
    seen = set()
    for m, ref, k, out, p in zip(t["contact_month"], t["dwelling_ref"], t["interview_no"], t["outcome"], t["psu"]):
        if out in (11, 12):
            seen.add((ref, month_index(m) - int(k) + 1, reg[p]))
    by = {}
    for _, _, r in seen:
        by[r] = by.get(r, 0) + 1
    by["all"] = len(seen)
    return by


def main(target, golden):
    import datetime as dt
    fails = []
    cells = read_cells(target)
    eff = fit_effects(cells)
    change = float(eff.mean())
    trigger = round(6.0 + round(change + 1e-9, 1), 1)
    regs, rates, nat = read_release(target)
    lab = lambda ym: dt.date(int(ym[:4]), int(ym[5:]), 1).strftime("%b-%Y")
    avg = np.mean([rates[lab("2026-%02d" % m)] for m in range(1, 13)], axis=0)
    designated = [regs[i] for i in range(len(regs)) if avg[i] >= trigger - 1e-9]
    q4 = q4_dwellings(target)
    print("verify: change %.4f (effects %s), restated trigger %.1f, designated %s" % (
        change, np.round(eff, 3).tolist(), trigger, designated))
    print("verify: fourth-quarter responding dwellings %s" % q4)
    from pypdf import PdfReader
    text = " ".join(" ".join((p.extract_text() or "").split()) for p in PdfReader(os.path.join(golden, PDF)).pages)
    names = {}
    with open(os.path.join(target, PSU), newline="") as f:
        for r in csv.DictReader(f):
            names[r["region_2025"]] = r["region_name"]
    for need in ("restated at %.1f per cent" % trigger, "%.1f percentage points higher" % round(change, 1),
                 format(q4["all"], ",")):
        if need not in text:
            fails.append("determination lacks %r" % need)
    for c in designated:
        if names[c] not in text:
            fails.append("determination lacks %s" % names[c])
    for c in [c for c in regs if c != "all"]:
        if format(q4[c], ",") not in text:
            fails.append("determination lacks %s's count" % names[c])
    rows = list(csv.DictReader(open(os.path.join(golden, CSVG), newline="")))
    if len(rows) != 36 * len(regs):
        fails.append("CSV rows %d" % len(rows))
    for r in rows:
        ym = r["month"]
        pub = rates[lab(ym)][regs.index(r["region"])]
        if abs(float(r["published_rate"]) - pub) > 1e-9:
            fails.append("CSV published %s %s" % (ym, r["region"]))
        # the note's section 4: before July 2025 the published rate plus the change as determined; from July 2025
        # as published (the transition months carried, not restated)
        want = round(pub + round(change + 1e-9, 1), 1) if ym < "2025-07" else pub
        if not re.fullmatch(r"\d+\.\d", r["redesigned_basis_rate"]) or abs(float(r["redesigned_basis_rate"]) - want) > 1e-9:
            fails.append("CSV redesigned %s %s %.3f against %.3f" % (ym, r["region"], float(r["redesigned_basis_rate"]),
                                                                     want))
    if trigger != 6.5 or designated != ["05", "07", "08", "09"]:
        fails.append("call: %.1f %s" % (trigger, designated))
    for f in fails[:12]:
        print("VERIFY FAIL  " + f)
    print("verify: %s (%d mismatches)" % ("PASS" if not fails else "FAIL", len(fails)))
    if fails:
        raise SystemExit(1)


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    task = os.path.dirname(here)
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(task, "target"),
         sys.argv[2] if len(sys.argv) > 2 else os.path.join(task, "golden"))
