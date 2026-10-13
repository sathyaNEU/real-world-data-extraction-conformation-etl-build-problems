"""The analysis, read from the shipped pack only: each economic region's monthly payroll jobs from January 2021 to
March 2027 (the yearly tables, the monthly file with the revision file's months in place of the preliminary ones,
municipalities to regions through the concordance), every seasonal adjustment route's test against the line, and the
two asks from the workplace returns."""
import csv
import os

import numpy as np
import openpyxl
import pandas as pd

import methods as M
import world as W

LINE = -1.0
FAMILY = ("reg12", "reg24", "reg12_2023", "ls_stl", "ls_stl13", "ls_classical", "pre_stl", "pre_classical",
          "splice_stl", "annual_ratio_stl", "q1_ratio_stl")          # the ranges the write-up states
GOLD = "reg12"                                                        # the bureau's own estimate in the certificate
DEFAULT = "stl"


def concordance(target):
    wb = openpyxl.load_workbook(os.path.join(target, W.FILES["conc"]), read_only=True)
    ws = wb["Municipalities"]
    rows = list(ws.iter_rows(values_only=True))
    head = rows[0]
    return {str(r[0]): str(r[3]) for r in rows[1:] if r[0]}


def wide(target):
    wb = openpyxl.load_workbook(os.path.join(target, W.FILES["wide"]), read_only=True)
    out = []
    for year in ("2021", "2022", "2023"):
        ws = wb[year]
        rows = list(ws.iter_rows(values_only=True))
        h = [i for i, r in enumerate(rows) if r and r[0] == "Municipality code"][0]
        for r in rows[h + 1:]:
            if not r or r[0] is None:
                continue
            for k in range(12):
                out.append(("%s-%02d" % (year, k + 1), str(r[0]), str(r[2]), int(r[3 + k])))
    return pd.DataFrame(out, columns=["month", "muni", "industry", "jobs"])


def long(target, revised=True):
    df = pd.read_csv(os.path.join(target, W.FILES["long"]), dtype={"municipality_code": str, "naics_sector": str})
    if revised:
        rv = pd.read_csv(os.path.join(target, W.FILES["rev"]), dtype={"municipality_code": str, "naics_sector": str})
        df = pd.concat([df[~df.reference_month.isin(set(rv.reference_month))], rv], ignore_index=True)
    return df.rename(columns={"reference_month": "month", "municipality_code": "muni", "naics_sector": "industry",
                              "payroll_jobs": "jobs"})[["month", "muni", "industry", "jobs"]]


def series(target, revised=True, drop_retired=False, from_2024=False):
    """months by economic region code, payroll jobs as the Agency's statistics place them"""
    conc = concordance(target)
    parts = [long(target, revised)]
    if not from_2024:
        parts.insert(0, wide(target))
    df = pd.concat(parts, ignore_index=True)
    if drop_retired:
        df = df[~df.muni.isin(W.AMALGAMATED)]
    df["region"] = df.muni.map(conc)
    assert df.region.notna().all()
    s = df.groupby(["month", "region"]).jobs.sum().unstack().sort_index()
    return s[W.RCODE]


def routes(s):
    """per cent by route and region"""
    F = M.Frame(list(s.index))
    out = {}
    for code in W.RCODE:
        out[code] = M.all_methods(F, s[code].values.astype(float))
    return pd.DataFrame(out).T


def designated(df, route):
    return [c for c in W.RCODE if df.loc[c, route] < LINE]


def sa_level(s, code, route=GOLD):
    """the bureau's estimate: the quarter's seasonally adjusted jobs and the 2026 average"""
    F = M.Frame(list(s.index))
    y = s[code].values.astype(float)
    lv = np.log(y)
    if route == "reg12":
        sf, _ = M.regression(F, lv)
    else:
        raise ValueError(route)
    sa = np.exp(lv - sf)
    return float(sa[F.q1].mean()), float(y[F.y26].mean())


# ------------------------------------------------------------------------------------------- the asks
def workplaces(target):
    df = pd.read_csv(os.path.join(target, W.FILES["work"]), dtype=str, keep_default_na=False)
    df["payroll_jobs"] = df.payroll_jobs.astype(int)
    conc = concordance(target)
    df["region"] = df.municipality_code.map(conc)
    return df


KEYS = {
    "workplace": lambda d: d.municipality_code + "|" + d.street_address,
    "account_number": lambda d: d.payroll_account + "|" + d.workplace_no,
    "business_address": lambda d: d.payroll_account.str[:9] + "|" + d.municipality_code + "|" + d.street_address,
    "legal_address": lambda d: d.legal_name + "|" + d.municipality_code + "|" + d.street_address,
}


def jobs_lost(df, key="workplace"):
    """by region: jobs lost between March 2026 and March 2027 at workplaces with none in March 2027 ('none') and at
    workplaces with fewer ('fewer')"""
    d = df.copy()
    d["k"] = KEYS[key](d)
    j = d.groupby(["region", "k", "reference_month"]).payroll_jobs.sum().unstack(fill_value=0)
    j = j.reindex(columns=["2026-03", "2027-03"], fill_value=0)
    lost = (j["2026-03"] - j["2027-03"]).clip(lower=0)
    none = j["2027-03"] == 0
    out = pd.DataFrame({"none": lost[none].groupby(level=0).sum(), "fewer": lost[~none].groupby(level=0).sum()})
    return out.reindex(W.RCODE).fillna(0).astype(int)


EMPLOYER_KEYS = {
    "business": lambda d: d.payroll_account.str[:9],
    "account": lambda d: d.payroll_account,
    "legal_name": lambda d: d.legal_name,
    "operating_name": lambda d: d.operating_name,
}


def employers(df, key="business"):
    d = df[(df.reference_month == "2027-03") & (df.payroll_jobs > 0)].copy()
    d["k"] = EMPLOYER_KEYS[key](d)
    return d.groupby("region").k.nunique().reindex(W.RCODE).fillna(0).astype(int)
