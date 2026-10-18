"""task245 checks: the three rules executed on the simulator's clean truth (R1 to R3c), the corridor of rebuild readings
and its margins (C1 to C3), every semantic fork row (SF1 to SF16), the Layer 1 mess, the separation of the call from the
asks, both asks' devices and stops, the pack's own registers, the input gates and the realism sweep."""
import datetime as dt
import os
import re

import numpy as np
import pandas as pd

import build as B
import estim as E
import params as P
import records as RC
import solve

N = [0]
GOLD = ["TBT1", "TH3", "TBT4"]
EASY = ["TH5", "TH6"]


def ok(cond, label):
    assert cond, label
    N[0] += 1


def _pdf_text(path):
    from pypdf import PdfReader
    return " ".join(" ".join(pg.extract_text().split()) for pg in PdfReader(path).pages)


def _xlsx_text(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    return " ".join(str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c is not None)


def _text(path):
    if path.endswith(".pdf"):
        return _pdf_text(path)
    if path.endswith(".xlsx"):
        return _xlsx_text(path)
    if path.endswith(".png"):
        return ""
    return open(path, encoding="utf-8").read()


def noted(R, key):
    return [P.CODES[i] for i in range(len(P.CODES)) if R["valid_days"][i] <= 35 and R[key][i] > 35]


def wind_keys(R):
    return [k for k in R if k.startswith("wind")]


def run(target, world, fig, golden):
    F = B.FILES
    W, X, cause = world["W"], world["X"], world["cause"]
    valid_t = cause == ""
    recs = RC.Records(world)
    # ---------------------------------------------------------------- the truth behind the golden
    Rt = E.readings(X, valid_t, W)
    A = solve.answer(target, F)
    R = A["R"]
    ok(fig["call"] == GOLD, "the golden call is the three plume stations")
    ok(E.call(Rt) == GOLD, "the clean truth gives the same call")
    true = (X.reshape(len(P.CODES), P.DAYS, 24).mean(2) > 50).sum(1)
    ok(all(true[P.CODES.index(c)] > 35 for c in GOLD) and all(true[i] <= 35 for i, c in enumerate(P.CODES) if c not in GOLD),
       "the simulated year: the three exceed and no other station does (%s)" % dict(zip(P.CODES, true.tolist())))
    # ---------------------------------------------------------------- R1: the raw material ships, nothing derived
    for k in ("council", "met", "register"):
        cols = open(os.path.join(target, F[k]), encoding="utf-8").readline().lower()
        ok(not re.search(r"whole|rebuilt|estimate|exceed|note|capture", cols), "R1 no derived column in %s" % F[k])
    ok(not re.search(r"(?i)whole.year|rebuil|exceed", _xlsx_text(os.path.join(target, F["operator"]))),
       "R1 the operator's workbook holds readings only")
    # ---------------------------------------------------------------- R2: the default path on the cut pack
    ok(noted(R, "pro_rata") == EASY, "R2 the count scaled by data capture notes Market Street and Westholme Avenue only")
    for m in ("linear", "lower", "higher", "nearest"):
        got = [P.CODES[i] for i in range(len(P.CODES)) if R["valid_days"][i] <= 35 and R["p904_" + m][i] > 50]
        ok(got and set(got) <= set(EASY), "R2 the 90.4th percentile (%s) notes only quiet-gap stations: %s" % (m, got))
    ok(all(R["valid_days"][i] <= 35 for i in range(len(P.CODES))), "R2 no station exceeds as measured")
    # ---------------------------------------------------------------- R3: the analyst test on the clean truth
    ok(noted(Rt, "pro_rata") == EASY, "R3 the easy route on the clean truth files Market Street and Westholme Avenue")
    ok(noted(Rt, E.REFERENCE) == GOLD, "R3 the decisive step alone (the wind rebuild) files the golden on the clean truth")
    for k in [k for k in Rt if k.startswith("net_")]:
        ok(noted(Rt, k) != GOLD, "R3b the network alone (%s) does not file the golden on the clean truth" % k)
    # ---------------------------------------------------------------- R3c: no plant, no method filed
    proc = _pdf_text(os.path.join(target, F["procedure"]))
    ok(not re.search(r"(?i)rebuil|estimat|interpolat|wind|scal(e|ing)|percentile|impute|infill", proc),
       "R3c the procedure files no way of estimating a lost day")
    ok("including any day on which the station has no valid daily mean" in proc, "R3c the estimand: every day of the year counts")
    ok("cannot be reported as meeting the objective" in proc and "35 or fewer" in proc, "R3c the note's definition is filed")
    for k in B.PACK:
        if F[k].endswith(".csv"):
            continue
        for sent in re.split(r"(?<=[.;])\s", _text(os.path.join(target, F[k]))):
            if re.search(r"(?i)\b(fault|stops?|stopped|alarm)\b", sent):
                ok(not re.search(r"(?i)\b(dust|plume|wind|loading|heavy)\b", sent),
                   "R3c no filed sentence ties a stop to the dust or the wind (%s): %s" % (k, sent[:90]))
    # ---------------------------------------------------------------- C1 to C3: the corridor of readings and its margins
    wk = wind_keys(R)
    ok(len(wk) == fig["readings"] and len(wk) >= 200, "C1 %d rebuild readings" % len(wk))
    for k in wk:
        ok(noted(R, k) == GOLD, "C1 every rebuild reading notes the three: %s %s" % (k, noted(R, k)))
    for c in GOLD:
        i = P.CODES.index(c)
        ok(min(R[k][i] for k in wk) >= 38, "C2 %s at least 38 under every rebuild (%d)" % (c, min(R[k][i] for k in wk)))
        ok(R["valid_days"][i] <= 31 and R["pro_rata"][i] <= 33.5, "C2 %s 31 or fewer over valid days, scaled at most 33.5" % c)
    for i, c in enumerate(P.CODES):
        if c not in GOLD:
            ok(max(R[k][i] for k in wk) <= 33, "C2 %s at most 33 under every rebuild (%d)" % (c, max(R[k][i] for k in wk)))
    for c in EASY:
        ok(R["pro_rata"][P.CODES.index(c)] >= 36.6, "C2 %s scaled at least 36.6, so every rounding of it is over 35" % c)
    for k in [k for k in R if k.startswith("net_") or k.startswith("interp")]:
        ok(noted(R, k) != GOLD, "C3 the rung %s does not land the call: %s" % (k, noted(R, k)))
    for k in [k for k in R if k.startswith("net_")]:
        ok(not any(c in noted(R, k) for c in GOLD), "C3 the network alone (%s) notes none of the three: %s" % (k, noted(R, k)))
    for k in [k for k in R if k.startswith("interp") or k.startswith("pro_rata")]:
        ok(not any(c in noted(R, k) for c in GOLD), "C3 interpolation in time or scaling (%s) notes none of the three: %s" % (k, noted(R, k)))
    # ---------------------------------------------------------------- semantic fork rows
    dm, vd = E.daily(A["xv"], A["valid"])
    ok(int(((dm > 49.4) & (dm < 50.6) & vd).sum()) == 0, "SF1 no valid daily mean within 0.6 of 50: above 50 or at least 50, same counts")
    ok(all(int((np.round(dm[i][vd[i]]) > 50).sum()) == R["valid_days"][i] for i in range(len(P.CODES))),
       "SF2 daily means rounded to whole units before the comparison: the same counts")
    for mode in ("as_gmt", "hour_end", "repeat_first"):
        xv2, va2 = grid_with(target, F, operator_alt(os.path.join(target, F["operator"]), mode))
        ok(E.call(E.readings(xv2, va2, A["W"])) == GOLD, "SF3/SF4/SF7 the operator's clock read %s: same call" % mode)
    ok("at least 18 of the day's 24 hourly means are valid" in proc, "SF5 the daily mean's validity is filed (section 3)")
    ok("The result covers every day of the year, including any day on which the station has no valid daily mean" in proc
       and "counts if the concentration at the station that day exceeded 50" in proc,
       "SF6 the lost days are in the count by that day's concentration (section 5.1)")
    ok(E.call(E.readings(A["xv"], A["valid"], solve.met(target, F, calm="vector"))) == GOLD,
       "SF8 calm and missing directions interpolated through the wind vector: same call")
    for c in P.CODES:
        i = P.CODES.index(c)
        lo, hi = min(R[k][i] for k in wk), max(R[k][i] for k in wk)
        ok(hi < 34 or lo > 36, "SF9 %s never within one of 35 under any rebuild" % c)
    ok(any(k.endswith("_pred") for k in wk) and any(k.endswith("_mix") for k in wk), "SF10 both partial-day handlings in the family")
    for thr in (0.0, 1.0):
        ok(E.call(E.readings(A["xv"], A["valid"], solve.met(target, F, wet_mm=thr))) == GOLD,
           "SF11 a wet day at more than %.1f mm: same call" % thr)
    ok(A["xv"].shape[1] == P.HOURS and P.DAYS == 365, "SF12 the calendar year 2027, 365 GMT days")
    for b in ("median", "council_bg", "rural", "two_way"):
        ok(any(("_%s_" % b) in k for k in wk), "SF13 the %s background in the family: same call" % b)
    cd = pd.read_csv(os.path.join(target, F["council"]), dtype=str, keep_default_na=False)
    ok(((cd["status"] != "R") == (cd["value"] == "")).all(), "SF14 every council hour not ratified as valid is blank")
    o = solve.operator(os.path.join(target, F["operator"]))
    ok(((o["status"] != "OK") == o["value"].isna()).all(), "SF15 every terminal hour not OK is blank")
    ev = [s for s in fig["stations"] if s["note"]]
    ok(all(s["lost_bearing"] / s["lost_days"] >= 2 * s["valid_bearing_share"] / 100 for s in ev),
       "SF16 each noted station's lost days had the terminal wind at least twice as often as its valid days")
    # ---------------------------------------------------------------- Layer 1: loud or neutral
    ok((cd["time_end"] == "24:00").sum() > 0 and not (cd["time_end"] == "00:00").any(), "L1 the council's hours end 01:00 to 24:00")
    ok(cd["date"].str.match(r"^\d\d/\d\d/2027$").all(), "L1 the council's dates are dd/mm/yyyy, all in 2027")
    raw = operator_alt(os.path.join(target, F["operator"]), "stamps")
    ok(raw.groupby("name").size().eq(P.HOURS).all(), "L1 the terminal's workbook carries every hour of 2027 once per station")
    # ---------------------------------------------------------------- separation: the asks sit outside the call
    pm = cd[cd["pollutant"] == "PM10"]
    ok(not pm.duplicated(subset=["site_code", "date", "time_end"]).any(), "S1 no PM10 hour repeated: the re-sent blocks are NO2 and NOx")
    ok(not any(c in GOLD for c in P.NO2_RESENT), "S1 no re-sent block at a plume station")
    ok(set(cd["pollutant"]) == {"PM10", "NO2", "NOx as NO2"}, "S1 the call reads the PM10 rows only (solve.grid)")
    # ---------------------------------------------------------------- ask A: NO2 annual means
    no2 = solve.no2_means(target, F)
    truth = recs.no2_truth()
    stops = {"export as it stands": solve.no2_means(target, F, fits="none"),
             "every hour rescaled": solve.no2_means(target, F, fits="all"),
             "re-sent block kept": solve.no2_means(target, F, resent="keep")}
    for c, v in no2.items():
        ok(abs(round(v, 1) - fig_station(fig, c)["no2"]) < 1e-9, "A golden NO2 %s" % c)
        ok(abs(round(v, 1) - v) <= 0.04, "A %s not within 0.01 of a rounding edge (%.3f)" % (c, v))
        ok(abs(v - truth[c]) < 0.03, "A %s recovers the simulated mean (%.3f against %.3f)" % (c, v, truth[c]))
    for k in ("export as it stands", "every hour rescaled"):
        for c, v in no2.items():
            ok(round(stops[k][c], 1) != round(v, 1), "A the stop '%s' misses %s (%.3f against %.3f)" % (k, c, stops[k][c], v))
    for c in P.NO2_RESENT:
        ok(round(stops["re-sent block kept"][c], 1) != round(no2[c], 1), "A the re-sent block kept moves %s (%.3f against %.3f)"
           % (c, stops["re-sent block kept"][c], no2[c]))
    log = _xlsx_text(os.path.join(target, F["servicelog"]))
    ok(log.count("converted on the logger at 1.88") == len(P.NO2), "A the referee: each new logger's conversion in the service log")
    ok("1.88" not in proc and "1.91" in proc, "A the procedure states 1.91 only")
    # ---------------------------------------------------------------- ask B: deposition rates
    dep = solve.deposition(target, F)
    dtruth = recs.deposit_truth()
    bstops = {"the laboratory's rates": solve.lab_rates(target, F),
              "the laboratory's area throughout": solve.deposition(target, F, area="lab"),
              "the Frisbee gauge's area throughout": solve.deposition(target, F, area="frisbee")}
    ended = solve.deposition(target, F, window="ended")
    void_days = solve.deposition(target, F, void="days")
    for name, v in dep.items():
        code = [s[0] for s in P.STATIONS if s[1] == name][0]
        ok(int(round(v)) == fig_station(fig, code)["dep"], "B golden deposition %s" % code)
        ok(abs(v - round(v)) <= 0.45, "B %s not on a rounding edge (%.2f)" % (code, v))
        ok(abs(v - dtruth[code]) < 0.3, "B %s recovers the simulated rate (%.2f against %.2f)" % (code, v, dtruth[code]))
        for k, vals in bstops.items():
            ok(round(vals[name]) != round(v), "B the stop '%s' misses %s (%.1f against %.1f)" % (k, code, vals[name], v))
    ok(sum(round(ended[n]) != round(dep[n]) for n in dep) >= 4, "B the hazard 'periods that ended in 2027' moves four gauges or more")
    ok(round(void_days["East Berth"]) != round(dep["East Berth"]), "B the hazard 'void period's days kept in the weights' moves TBT2")
    plan = _pdf_text(os.path.join(target, F["plan"]))
    ok("227 mm" in plan and "Frisbee" in plan and "305" not in plan, "B the plan names the Frisbee gauge only")
    ok(set(solve.gauge_swaps(target, F).values()) == {P.GAUGE_SWITCH}, "B the referee: every gauge swapped at the June exchange")
    ok(_xlsx_text(os.path.join(target, F["deposit"])).count("0.0731") >= 60, "B the laboratory keeps the funnel gauge's area")
    # ---------------------------------------------------------------- the pack's own registers (H9)
    readme = open(os.path.join(target, F["readme"]), encoding="utf-8").read()
    dic = open(os.path.join(target, F["dict"]), encoding="utf-8").read()
    ok({f for f in os.listdir(target)} == {F[k] for k in B.PACK}, "H9 the target holds the pack and nothing else")
    for k in B.PACK:
        if k != "readme":
            ok(F[k] in readme, "H9 README lists %s" % F[k])
    for k in ("council", "operator", "met", "register", "servicelog", "deposit", "asr2026", "bgmap"):
        ok(F[k] in dic, "H9 the dictionary covers %s" % F[k])
    # ---------------------------------------------------------------- input gates and realism
    files = [F[k] for k in B.PACK]
    ok(len(files) >= 10, "gate: %d files" % len(files))
    ok(len({f.rsplit(".", 1)[1] for f in files}) >= 3, "gate: three formats or more")
    ok(sum(1 for _ in open(os.path.join(target, F["council"]))) > 25000, "gate: a file of 25,000 rows or more")
    ok(len(B.DISTRACTORS) >= 2, "gate: two distractors")
    for f in files + sorted(os.listdir(golden)):
        p = os.path.join(target, f) if f in files else os.path.join(golden, f)
        t = _text(p)
        ok("\u2014" not in t and "\u2013" not in t, "realism: no long dash in %s" % f)
        ok(not re.search(r"(?i)fictional|fictitious|this exercise|constructed for|synthetic", t), "realism: no build words in %s" % f)
    return N[0]


def fig_station(fig, code):
    return [s for s in fig["stations"] if s["code"] == code][0]


def operator_alt(path, mode):
    """the terminal's workbook with its clock read another way: 'as_gmt' (UK clock taken as GMT), 'hour_end' (the stamp
    taken as the end of the hour), 'repeat_first' (the repeated 01:00 of 31 October read once, as GMT), 'stamps' (the
    golden reading, one row per station and hour)"""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    rows = []
    for ws in wb.worksheets:
        head = None
        for r in ws.iter_rows(values_only=True):
            if r and r[0] == "Date/time":
                head = list(r)
                continue
            if head is None or r[0] is None:
                continue
            t = dt.datetime.strptime(r[0], "%d/%m/%Y %H:%M")
            for k in range(1, len(head), 2):
                rows.append((head[k].replace(" PM10", ""), t, r[k], r[k + 1]))
    df = pd.DataFrame(rows, columns=["name", "t", "value", "status"])
    df["n"] = df.groupby(["name", "t"]).cumcount()
    bst = (df["t"] >= pd.Timestamp("2027-03-28 02:00")) & (df["t"] < pd.Timestamp("2027-10-31 01:00"))
    rep = (df["t"] == pd.Timestamp("2027-10-31 01:00")) & (df["n"] == 0)
    if mode in ("stamps", "hour_end"):
        gmt = df["t"] - pd.to_timedelta((bst | rep).astype(int), unit="h")
    elif mode == "as_gmt":
        gmt = df["t"]
    else:
        gmt = df["t"] - pd.to_timedelta(bst.astype(int), unit="h")
    df["h"] = ((gmt - pd.Timestamp(P.START)) / pd.Timedelta(hours=1)).astype(int) - (1 if mode == "hour_end" else 0)
    df = df[(df["h"] >= 0) & (df["h"] < P.HOURS)]
    return df.drop_duplicates(subset=["name", "h"], keep="first")


def grid_with(target, F, op):
    """solve.grid with the terminal's hours placed by another reading of its clock"""
    codes, xv, valid = solve.grid(target, F)
    reg = pd.read_csv(os.path.join(target, F["register"]), dtype=str)
    names = dict(zip(reg["site_name"], reg["site_code"]))
    for name, g in op.groupby("name"):
        i = codes.index(names[name])
        xv[i] = 0.0
        valid[i] = False
        good = (g["status"] == "OK").values
        h = g["h"].values[good]
        xv[i, h] = g["value"].values[good].astype(float)
        valid[i, h] = True
    return xv, valid
