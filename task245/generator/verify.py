"""task245 independent verifier: reads only the shipped bundle (target/), the goldens, figures.json and submission.md;
places the terminal's hours by the Europe/London time zone (not solve.py's BST window), builds its own grid, mast and
reference rebuild (the network median plus the station's mean increment by wind sector, speed band and wet day), recomputes
the call and every station's whole-year count, both asks by its own code, reads back every format the pack ships, and
requires the goldens and the write-up to carry the same figures."""
import csv
import datetime as dt
import json
import math
import os
import re
import sys
import warnings
from zoneinfo import ZoneInfo

import numpy as np
from openpyxl import load_workbook

T0 = dt.datetime(2027, 1, 1, tzinfo=dt.timezone.utc)
H, D = 8760, 365
LONDON = ZoneInfo("Europe/London")


def pdf_text(p):
    from pypdf import PdfReader
    return " ".join(" ".join(pg.extract_text().split()) for pg in PdfReader(p).pages)


def rows_of(path, header_first):
    """a workbook's rows after the header row whose first cell is header_first, as dicts"""
    out = []
    for ws in load_workbook(path, data_only=True).worksheets:
        head = None
        for r in ws.iter_rows(values_only=True):
            if r and r[0] == header_first:
                head = list(r)
                continue
            if head and r[0] is not None:
                out.append(dict(zip(head, r)))
    return out


def hour_of(end_utc):
    return int((end_utc - T0).total_seconds() // 3600) - 1


def council(path):
    """PM10 and NO2 rows of the council's export, each with its hour index"""
    out = []
    with open(path, encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            d = dt.datetime.strptime(r["date"], "%d/%m/%Y").replace(tzinfo=dt.timezone.utc)
            hh, mm = map(int, r["time_end"].split(":"))
            r["h"] = hour_of(d + dt.timedelta(hours=hh, minutes=mm))
            out.append(r)
    return out


def operator(path):
    """the terminal's rows: (station name, hour index, value, status), the clock read through Europe/London, the second
    of two equal stamps taken as the later instant"""
    seen = {}
    out = []
    for r in rows_of(path, "Date/time"):
        local = dt.datetime.strptime(r["Date/time"], "%d/%m/%Y %H:%M")
        names = [k[:-5] for k in r if k and k.endswith(" PM10")]
        fold = seen.get(local, 0)
        seen[local] = fold + 1
        start = local.replace(tzinfo=LONDON, fold=fold).astimezone(dt.timezone.utc)
        h = int((start - T0).total_seconds() // 3600)
        for n in names:
            out.append((n, h, r[n + " PM10"], r[n + " status"]))
    return out


def mast(path):
    theta = np.full(H, np.nan)
    speed = np.full(H, np.nan)
    rain = np.zeros(H)
    with open(path, encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            h = hour_of(dt.datetime.strptime(r["datetime_end_gmt"], "%Y-%m-%dT%H:%MZ").replace(tzinfo=dt.timezone.utc))
            if r["wind_dir_deg"]:
                theta[h] = float(r["wind_dir_deg"])
            if r["wind_speed_ms"]:
                speed[h] = float(r["wind_speed_ms"])
            if r["rain_mm"]:
                rain[h] = float(r["rain_mm"])
    have = np.flatnonzero(~np.isnan(theta))
    full = theta.copy()
    for h in np.flatnonzero(np.isnan(theta)):
        j = np.searchsorted(have, h)
        cands = [have[k] for k in (j - 1, j) if 0 <= k < len(have)]
        full[h] = theta[min(cands, key=lambda c: (abs(c - h), c))]
    sp = np.flatnonzero(~np.isnan(speed))
    speed = np.interp(np.arange(H), sp, speed[sp])
    wet = np.repeat(rain.reshape(D, 24).sum(1) > 0.2, 24)
    return full % 360, speed, wet


def reference_count(x, i, theta, speed, wet):
    """the station's whole-year count: valid daily means as measured, the other days from their valid hours and the rebuilt
    ones (median of the other stations that hour plus the station's mean increment in the hour's wind cell)"""
    others = np.delete(x, i, axis=0)
    with np.errstate(all="ignore"):
        bg = np.nanmedian(others, 0)
    inc = x[i] - bg
    sec = (((theta + 11.25) % 360) // 22.5).astype(int)
    band = (speed >= 4.0).astype(int)
    cell = (sec * 4 + band) * 2 + wet.astype(int)
    use = ~np.isnan(inc)
    overall = inc[use].mean()
    cmean, smean = {}, {}
    for c in np.unique(cell[use]):
        sel = use & (cell == c)
        if sel.sum() >= 5:
            cmean[c] = inc[sel].mean()
    for s in np.unique(sec[use]):
        smean[s] = inc[use & (sec == s)].mean()
    pred = bg + np.array([cmean.get(c, smean.get(s, overall)) for c, s in zip(cell, sec)])
    hrs = np.where(np.isnan(x[i]), pred, x[i]).reshape(D, 24)
    nvalid = (~np.isnan(x[i])).reshape(D, 24).sum(1)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", RuntimeWarning)
        measured = np.nanmean(x[i].reshape(D, 24), 1)
    day = np.where(nvalid >= 18, measured, hrs.mean(1))
    return int((day > 50).sum()), int((measured[nvalid >= 18] > 50).sum()), int((nvalid >= 18).sum())


def main(target, golden, figures, submission):
    fig = json.load(open(figures))
    P = lambda f: os.path.join(target, f)
    reg = list(csv.DictReader(open(P("thornmouth_aq_station_register.csv"), encoding="utf-8")))
    codes = [r["site_code"] for r in reg]
    code_of = {r["site_name"]: r["site_code"] for r in reg}
    proc = pdf_text(P("tbc_aq_network_qaqc_procedure_v4.pdf"))
    assert "at least 18 of the day's 24 hourly means are valid" in proc, "verify: procedure, daily means"
    assert "such a day counts if the concentration at the station that day exceeded 50" in proc, "verify: procedure, 5.1"
    assert "1 ppb = 1.91" in proc, "verify: procedure NO2 basis"
    eml = open(P("aq_audit_2027_pm10_note_request.eml"), encoding="utf-8").read()
    assert eml.startswith("From: ") and "Subject: " in eml, "verify: request"
    assert "TABLE A.3" in pdf_text(P("asr2027_table_a3_no2_2026.pdf")).upper(), "verify: last year's table"
    # the grid
    x = np.full((len(codes), H), np.nan)
    cd = council(P("thornmouth_council_aq_hourly_2027.csv"))
    for r in cd:
        if r["pollutant"] == "PM10" and r["status"] == "R":
            x[codes.index(r["site_code"]), r["h"]] = float(r["value"])
    op = operator(P("tbt_pm10_permit_monitoring_2027.xlsx"))
    for name, h, v, st in op:
        if st == "OK" and 0 <= h < H:
            x[codes.index(code_of[name]), h] = float(v)
    theta, speed, wet = mast(P("port_met_mast_hourly_2027.csv"))
    counts = {c: reference_count(x, i, theta, speed, wet) for i, c in enumerate(codes)}
    noted = [c for c in codes if counts[c][1] <= 35 and counts[c][0] > 35]
    assert noted == fig["call"], ("verify: the call", noted)
    S = {s["code"]: s for s in fig["stations"]}
    for c in codes:
        year, valid_count, valid_days = counts[c]
        assert year == S[c]["year_count"], ("verify: whole-year count", c, year, S[c]["year_count"])
        assert valid_count == S[c]["valid_count"] and valid_days == S[c]["valid_days"], ("verify: valid days", c)
        assert S[c]["interval"][0] <= year <= S[c]["interval"][1], ("verify: interval", c)
        assert (S[c]["interval"][0] > 35) == (c in noted) and (S[c]["interval"][1] <= 35) == (c not in noted), ("range", c)
    # ask A: NO2, each hour once, the new loggers' hours rescaled to 1.91
    log = rows_of(P("aq_network_service_log_2027.xlsx"), "Date")
    fits = {}
    for r in log:
        if r["Type"] == "Logger":
            local = dt.datetime.strptime(r["Date"] + " " + r["Time (UK clock)"], "%d/%m/%Y %H:%M")
            start = local.replace(tzinfo=LONDON).astimezone(dt.timezone.utc)
            fits[r["Site"]] = (int((start - T0).total_seconds() // 3600), float(re.search(r"at ([\d.]+) ug/m3", r["Notes"]).group(1)))
    seen, vals = set(), {}
    for r in cd:
        if r["pollutant"] != "NO2" or r["status"] != "R":
            continue
        key = (r["site_code"], r["h"])
        if key in seen:
            continue
        seen.add(key)
        h0, fac = fits[r["site_code"]]
        v = float(r["value"]) * (1.91 / fac if r["h"] >= h0 else 1.0)
        vals.setdefault(r["site_code"], []).append(v)
    for c, v in vals.items():
        assert round(sum(v) / len(v), 1) == S[c]["no2"], ("verify: NO2", c, sum(v) / len(v))
    # ask B: deposition, each period over the collecting area of the gauge exposed then
    plan = pdf_text(P("tbt_dust_management_monitoring_plan_rev3.pdf"))
    dia = float(re.search(r"Frisbee-type deposit gauge, (\d+) mm collecting diameter", plan).group(1))
    swap = {r["Site"]: dt.datetime.strptime(r["Date"], "%d/%m/%Y").date() for r in log
            if r["Type"] == "Deposit gauge" and "Frisbee" in r["Notes"]}
    agg = {}
    for r in rows_of(P("tbt_deposit_gauge_results_2027.xlsx"), "Gauge"):
        on = dt.datetime.strptime(r["Date on"], "%d/%m/%Y").date()
        if on.year != 2027 or str(r["Remarks"] or "").startswith("VOID"):
            continue
        c = code_of[r["Site"]]
        area = math.pi * (dia / 2000.0) ** 2 if on >= swap[c] else float(r["Collecting area (m2)"])
        a = agg.setdefault(c, [0.0, 0])
        a[0] += float(r["Total deposit (mg)"]) / area
        a[1] += int(r["Days exposed"])
    for c, (num, den) in agg.items():
        assert int(round(num / den)) == S[c]["dep"], ("verify: deposition", c, num / den)
    # the goldens carry the figures
    wb = load_workbook(os.path.join(golden, "thornmouth_pm10_2027_data_capture_review.xlsx"), data_only=True)
    got = 0
    for row in wb["Review"].iter_rows(min_row=5, values_only=True):
        if row[0] in S:
            s = S[row[0]]
            assert row[3] == s["valid_days"] and row[4] == s["valid_count"] and row[5] == s["year_count"], ("wb", row)
            assert row[6] == "%d to %d" % tuple(s["interval"]) and (row[7] == "Note") == s["note"], ("wb", row)
            assert row[8] == s["no2"] and row[9] == s["dep"], ("wb asks", row)
            got += 1
    assert got == 12, "verify: workbook rows"
    note = pdf_text(os.path.join(golden, "asr2027_pm10_data_capture_note.pdf"))
    from pypdf import PdfReader
    assert len(PdfReader(os.path.join(golden, "asr2027_pm10_data_capture_note.pdf")).pages) == 1, "verify: one page"
    for c in noted:
        assert "(%s)" % c in note, ("verify: note names", c)
    for c in codes:
        if c not in noted:
            assert "(%s)" % c not in note.split("The days these")[0], ("verify: the note's opening names the four only", c)
    sub = open(submission, encoding="utf-8").read()
    for c in codes:
        s = S[c]
        assert re.search(r"- %s \(%s\): %d over the whole year \(%d to %d\); %s" % (
            re.escape(s["name"]), c, s["year_count"], s["interval"][0], s["interval"][1], "note" if s["note"] else "no note"), sub), \
            ("verify: submission row", c)
        if s["no2"] is not None:
            assert "- %s (%s) %.1f" % (s["name"], c, s["no2"]) in sub, ("verify: submission NO2", c)
        if s["dep"] is not None:
            assert "- %s (%s) %d" % (s["name"], c, s["dep"]) in sub, ("verify: submission deposition", c)
    print("verify: call (%s), every whole-year count equal to the golden and inside its range, both asks, every format "
          "read back, goldens and write-up agree" % ", ".join(noted))


if __name__ == "__main__":
    task = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    main(os.path.join(task, "target"), os.path.join(task, "golden"),
         os.path.join(task, "generator", "figures.json"), os.path.join(task, "submission.md"))
