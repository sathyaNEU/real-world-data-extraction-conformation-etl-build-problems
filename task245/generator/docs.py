"""task245 documents: the air quality officer's request (eml), the council's QA/QC procedure for the network (pdf), the
terminal operator's dust management and monitoring plan (pdf), the 2026 report's NO2 table (pdf), the operator's PM10
workbook, the service log and the deposit-gauge results (xlsx), the data dictionary and the readme (md)."""
import datetime as dt
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess

COUNCIL = "Thornmouth Borough Council"
OPERATOR = "Thornmouth Bulk Terminal Ltd"
CONTRACTOR = "Ashcroft Air Quality Services Ltd"
PEOPLE = {"aqo": ("Harriet Hopkins", "harriet.hopkins", "thornmouth.gov.uk"),
          "analyst": ("Conor Robinson", "conor.robinson", "thornmouth.gov.uk"),
          "auditor": ("Marc Bevan", "marc.bevan", "ashcroft-aq.co.uk"),
          "operator": ("Jemma Holloway", "jemma.holloway", "thornmouthbulk.co.uk")}


def _addr(k):
    name, user, domain = PEOPLE[k]
    return "%s <%s@%s>" % (name, user, domain)


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.6, leading=13.4, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=13.5, leading=17, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=14,
                        spaceBefore=8, spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10.5, textColor=colors.HexColor("#444444"))
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=8.4, leading=10.2)
    return body, h1, h2, small, cell


def _grid(rows, widths, cell):
    t = Table([[Paragraph(str(c), cell) for c in r] for r in rows], colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA0AB")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E3E8EF"))]))
    return t


def _footer(text):
    def foot(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(20 * mm, 10 * mm, text)
        canvas.drawRightString(A4[0] - 20 * mm, 10 * mm, "Page %d" % d.page)
        canvas.restoreState()
    return foot


def _doc(path, title, author, subject):
    return SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                             bottomMargin=18 * mm, title=title, author=author, creator=author, subject=subject,
                             invariant=1)


# ------------------------------------------------------------------ the request (eml)
def request(path, F):
    body = f"""From: {_addr("aqo")}
To: {_addr("analyst")}
Cc: {_addr("auditor")}
Date: Fri, 17 Mar 2028 09:12:44 +0000
Subject: 2027 data audit: PM10 data-capture note before the ASR
MIME-Version: 1.0
Content-Type: text/plain; charset="utf-8"
Content-Transfer-Encoding: 8bit

Conor,

Marc's audit of the network's 2027 data closes on Friday 31 March. Before he signs off the year's statistics for
the Annual Status Report he wants us to settle which PM10 stations' 2027 results go in with the data-capture note
(section 5 of the QA/QC procedure). That covers all twelve stations, ours and the five the terminal runs under its
permit and shares with us.

Could you put together the review workbook ({F["deliv_xlsx"]}) with a row for every station, and a one-page note
for Marc ({F["deliv_pdf"]}) with a chart he can take straight into the audit file. Marc also checks two other
figures for each site at the audit, so the workbook should carry them beside the PM10 results: the 2027 NO2
annual mean at our six NO2 stations, and the 2027 dust deposition rate at the terminal's five gauges.

In the folder:
- our 2027 export from the network database (ratified, all seven of our stations);
- the terminal's 2027 PM10 workbook, which Jemma sent on 6 March in their logger's own layout;
- the port's met mast for the year;
- the station register;
- both technicians' service log for the year;
- the terminal's deposit gauge results from their laboratory, and their dust monitoring plan (revision 3);
- the procedure itself (version 4);
- last year's ASR NO2 table and the background map extract we used for it.

A draft by Wednesday 29th would give Marc a day with it.

Thanks,

Harriet Hopkins
Air Quality Officer, Environmental Protection
{COUNCIL}
"""
    with open(path, "w") as fh:
        fh.write(body)


# ------------------------------------------------------------------ the council's QA/QC procedure (pdf)
def procedure(path, F):
    body, h1, h2, small, cell = _styles()
    doc = _doc(path, "Air quality monitoring network: QA/QC procedure AQ-QP-04", COUNCIL,
               "Data handling, ratification and annual statistics for the borough's air quality monitoring network")
    s = [Paragraph("%s, Environmental Protection" % COUNCIL, small),
         Paragraph("Air quality monitoring network: QA/QC procedure AQ-QP-04", h1),
         _grid([["Reference", "Version", "Issued", "Owner", "Approved by", "Next review"],
                ["AQ-QP-04", "4", "11 January 2027", "Air Quality Officer", "Head of Environmental Protection",
                 "January 2029"]], [22 * mm, 16 * mm, 28 * mm, 30 * mm, 42 * mm, 26 * mm], cell),
         Spacer(1, 4)]
    s += [Paragraph("1. Scope", h2),
          Paragraph("This procedure covers the automatic monitoring stations of the borough's air quality network: the "
                    "council's seven stations, and the five PM10 stations that %s runs around the bulk terminal under its "
                    "environmental permit and shares with the council under the data-sharing agreement of 2025 (Annex "
                    "A). It sets out how hourly data are handled, how they are ratified, and how the annual statistics "
                    "in the council's Annual Status Report (ASR) are worked out. The council's QA/QC contractor audits "
                    "each year's data and statistics before the ASR is submitted." % OPERATOR, body)]
    s += [Paragraph("2. Hourly data", h2),
          Paragraph("2.1 The network works in GMT throughout the year. Each hourly mean is stamped with the end of the "
                    "hour it covers, so the hour ending 01:00 is the first hour of a day and the hour ending 24:00 its "
                    "last.", body),
          Paragraph("2.2 The council's stations report to the network database directly. The terminal's stations are "
                    "logged on the operator's own system, ratified on the terms of this procedure and sent to the "
                    "council each month in the operator's workbook, each hour with its status (OK, FAULT, REJECTED, "
                    "SERVICE, POWER or OFF, as R, F, X, S, P and O below); they enter the database after conversion to "
                    "the network's time base.", body),
          Paragraph("2.3 Ratified hourly data carry a status code. R: ratified and valid. F: analyser fault (no "
                    "reading). X: rejected at ratification. S: service or calibration. P: power or telemetry lost. O: "
                    "station out of service. Only R hours are used in any statistic; data lost or rejected are not "
                    "replaced in the ratified record.", body),
          Paragraph("2.4 When a TEOM stops on a fault, its readings earlier that day are rejected as unstable. After a "
                    "restart the analyser's readings are rejected to the end of that day while the mass baseline "
                    "settles.", body)]
    s += [Paragraph("3. Daily means", h2),
          Paragraph("A daily mean is the mean of the day's valid hourly means (00:00 to 24:00 GMT). It is valid where "
                    "at least 18 of the day's 24 hourly means are valid.", body)]
    s += [Paragraph("4. Objectives", h2),
          _grid([["Pollutant", "Objective", "Measured as"],
                 ["PM10", "50 µg/m³, not to be exceeded more than 35 times a year", "24-hour mean"],
                 ["PM10", "40 µg/m³", "Annual mean"],
                 ["NO2", "200 µg/m³, not to be exceeded more than 18 times a year", "1-hour mean"],
                 ["NO2", "40 µg/m³", "Annual mean"]], [22 * mm, 90 * mm, 40 * mm], cell),
          Spacer(1, 3),
          Paragraph("Source: the Air Quality (England) Regulations 2000, as amended.", small)]
    s += [Paragraph("5. Annual statistics for the ASR", h2),
          Paragraph("5.1 A station's result against the PM10 24-hour objective is the number of days in the calendar "
                    "year on which the daily mean concentration at the station exceeded 50 µg/m³. The result "
                    "covers every day of the year, including any day on which the station has no valid daily mean: such "
                    "a day counts if the concentration at the station that day exceeded 50 µg/m³, and not "
                    "otherwise.", body),
          Paragraph("5.2 The ASR also gives, for each station, the number of days above 50 µg/m³ among its "
                    "valid daily means. Where that count is 35 or fewer but the station's result for the year is more "
                    "than 35, the station's result carries the data-capture note: the station cannot be reported as "
                    "meeting the objective.", body),
          Paragraph("5.3 NO2 is reported in µg/m³ at 20 °C and 101.3 kPa (1 ppb = 1.91 "
                    "µg/m³). A station's NO2 annual mean is the mean of its valid hourly means in the calendar "
                    "year.", body),
          Paragraph("5.4 Data capture is the share of the year's days with a valid daily mean (PM10) or of the year's "
                    "hours with a valid hourly mean (NO2).", body)]
    s += [Paragraph("6. Audit", h2),
          Paragraph("The QA/QC contractor audits the ratified data, the service records and the annual statistics "
                    "after the year closes, and signs off the network's statistics before the ASR is submitted.", body)]
    s += [Paragraph("7. Checks and calibration", h2),
          Paragraph("7.1 NO2 analysers: zero and span check every two weeks against a certified cylinder; a response "
                    "outside 5 per cent of the cylinder is investigated before the data are ratified.", body),
          Paragraph("7.2 PM10 analysers: inlet cleaning, flow and leak checks at each quarterly service; TEOM filters "
                    "exchanged on each visit and BAM tape and nozzle checked; flow audits by the QA/QC contractor twice "
                    "a year.", body),
          Paragraph("7.3 Visits, faults, services and calibrations are entered in the service log by the technician "
                    "who carries them out.", body)]
    reg = [["Site", "Name", "Operator", "Site type", "PM10 analyser", "NO2"]]
    for st in P.STATIONS:
        reg.append([st[0], st[1], "Council" if st[2] == "TBC" else "Terminal", st[3], st[8],
                    "yes" if st[0] in P.NO2 else ""])
    s += [Paragraph("Annex A. Network stations", h2),
          _grid(reg, [14 * mm, 34 * mm, 20 * mm, 32 * mm, 30 * mm, 12 * mm], cell)]
    s += [Paragraph("Document history", h2),
          _grid([["Version", "Issued", "Change"],
                 ["1", "April 2019", "First issue"],
                 ["2", "June 2022", "Ratification by the QA/QC contractor; status codes"],
                 ["3", "March 2025", "The terminal's stations brought into the network (data-sharing agreement)"],
                 ["4", "January 2027", "Section 2.4 added after the 2026 audit; section 7"]],
                [18 * mm, 28 * mm, 110 * mm], cell)]
    doc.build(s, onFirstPage=_footer("AQ-QP-04 v4"), onLaterPages=_footer("AQ-QP-04 v4"))
    postprocess.scrub(path, COUNCIL, "2027-01-11")


# ------------------------------------------------------------------ the operator's plan (pdf)
def plan(path, F):
    body, h1, h2, small, cell = _styles()
    doc = _doc(path, "Dust management and monitoring plan, revision 3", OPERATOR,
               "Dust controls and ambient monitoring at the Thornmouth Bulk Terminal")
    s = [Paragraph(OPERATOR, small),
         Paragraph("Dust management and monitoring plan", h1),
         Paragraph("Revision 3, May 2027. Prepared under condition 3.4 of the terminal's environmental permit.", small),
         Spacer(1, 4)]
    s += [Paragraph("1. The terminal", h2),
          Paragraph("The terminal imports iron ore, coal and aggregates through its deep-water berth. Cargo is "
                    "discharged by grab cranes onto the conveyor and stacked on the open stockyard, about 18 hectares "
                    "centred on national grid reference E %d, N %d; it leaves by rail and road. Water sprays run on the "
                    "conveyor transfer points and on the stockpiles in dry weather." % P.STOCKYARD_OS, body)]
    s += [Paragraph("2. Ambient PM10 monitoring", h2),
          _grid([["Station", "Position", "Analyser"]] +
                [[st[1], {"TBT1": "north boundary, by the river quay", "TBT2": "east of the stockyard, at the berth",
                          "TBT3": "rail sidings, north-east", "TBT4": "south boundary", "TBT5": "west gate, north-west"}[st[0]],
                  st[8]] for st in P.STATIONS if st[2] == "TBT"], [32 * mm, 70 * mm, 40 * mm], cell),
          Spacer(1, 4),
          Paragraph("The stations log hourly means on the terminal's data system, in UK clock time, each hour stamped "
                    "with its start. The year's data go to the council each month as a workbook, one sheet per quarter, "
                    "with a status column for each station.", body),
          Paragraph("The terminal's technician visits every station on Mondays and Thursdays from 10:00. A stopped "
                    "analyser is restarted, and its filter exchanged, on the next round.", body)]
    s += [Paragraph("3. Dust deposition", h2),
          Paragraph("A deposit gauge stands at each PM10 station. Gauges are exchanged on the first Monday of each "
                    "month and the bottles analysed by the terminal's contracted laboratory for insoluble and soluble "
                    "deposit; the laboratory reports each exposure period's total deposit and deposition rate.", body),
          Paragraph("A period's deposition rate is its total deposit (insoluble plus soluble) divided by the gauge's "
                    "collecting area and by the days the gauge was exposed. A gauge's deposition rate for a year is the "
                    "mean of the rates of the exposure periods that started in that year, each weighted by its days. "
                    "Periods the laboratory reports void are left out.", body),
          Paragraph("The terminal works to a site action level of 200 mg/m² per day for any single exposure "
                    "period, at which the dust controls are reviewed.", body)]
    s += [Paragraph("4. Reporting", h2),
          Paragraph("The terminal sends the council its PM10 workbook each month and its deposit gauge results each "
                    "quarter, and reports the year's results to the Environment Agency in its annual permit return.", body)]
    s += [Paragraph("Annex A. Monitoring equipment", h2),
          _grid([["Item", "Equipment", "Notes"],
                 ["PM10 analysers", "TEOM 1405 (North Gate, East Berth, South Yard, West Gate); Fidas 200 (Rail Sidings)",
                  "Inlet 3.0 m above ground; filter exchanged on every round"],
                 ["Deposit gauges", "Frisbee-type deposit gauge, 227 mm collecting diameter, on a 1.5 m post",
                  "Polyethylene collecting bottle; 20 ml of algicide added at each exchange"],
                 ["Data system", "Terminal environmental data system", "Hourly means logged in UK clock time"]],
                [30 * mm, 74 * mm, 50 * mm], cell)]
    s += [Paragraph("Revision history", h2),
          _grid([["Rev.", "Date", "Change"],
                 ["1", "June 2019", "First issue"],
                 ["2", "October 2025", "Monitoring data shared with the council (data-sharing agreement)"],
                 ["3", "May 2027", "Annex A updated; technician rounds on Mondays and Thursdays"]],
                [14 * mm, 28 * mm, 110 * mm], cell)]
    doc.build(s, onFirstPage=_footer("TBT DMMP rev. 3"), onLaterPages=_footer("TBT DMMP rev. 3"))
    postprocess.scrub(path, OPERATOR, "2027-05-24")


# ------------------------------------------------------------------ last year's ASR NO2 table (pdf)
def asr2026(path, F, R):
    body, h1, h2, small, cell = _styles()
    doc = _doc(path, "Annual Status Report 2027: NO2 at the automatic stations", COUNCIL,
               "2026 annual mean NO2 at the council's automatic monitoring stations")
    rows = [["Site", "Site type", "Data capture 2026 (%)", "Annual mean NO2 2026 (µg/m³)"]]
    for code, (cap, mean) in R.no2_2026().items():
        rows.append(["%s (%s)" % (P.NAME[code], code), [s[3] for s in P.STATIONS if s[0] == code][0], "%.1f" % cap, "%.1f" % mean])
    s = [Paragraph(COUNCIL, small),
         Paragraph("2027 Air Quality Annual Status Report (2026 data): Table A.3", h1),
         Paragraph("Annual mean NO2 at the automatic monitoring stations, 2026.", body),
         _grid(rows, [52 * mm, 34 * mm, 34 * mm, 40 * mm], cell),
         Spacer(1, 4),
         Paragraph("Notes: concentrations at 20 °C and 101.3 kPa; data ratified by the council's QA/QC contractor. "
                   "No station had data capture below 75 per cent, so no annual mean has been annualised.", small)]
    doc.build(s, onFirstPage=_footer("ASR 2027, Table A.3"), onLaterPages=_footer("ASR 2027, Table A.3"))
    postprocess.scrub(path, COUNCIL, "2027-06-23")


# ------------------------------------------------------------------ workbooks
def _sheet_widths(ws, widths):
    for k, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(k)].width = w


def operator_workbook(path, R):
    wb = Workbook()
    wb.remove(wb.active)
    for title, df in R.operator_frames():
        ws = wb.create_sheet(title)
        ws["A1"] = "%s: ambient PM10 monitoring, %s" % (OPERATOR, title)
        ws["A1"].font = Font(bold=True)
        ws["A2"] = ("Hourly means in ug/m3. Time is UK clock time and marks the start of the hour. Status as ratified: OK, "
                    "FAULT, REJECTED, SERVICE, POWER, OFF.")
        ws.append([])
        ws.append(list(df.columns))
        for c in ws[4]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDE3EA")
        for row in df.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and v != v) else v for v in row])
        ws.freeze_panes = "B5"
        _sheet_widths(ws, [17] + [13, 10] * 5)
    wb.save(path)
    postprocess.scrub(path, OPERATOR, "2028-03-06")
    postprocess.finish_office(path, "2027-04-03T07:55:00Z", "2028-03-06T14:20:00Z", OPERATOR)


def table_workbook(path, df, sheet, title, note, producer, stamp, widths, created):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws["A1"] = title
    ws["A1"].font = Font(bold=True)
    ws["A2"] = note
    ws.append([])
    ws.append(list(df.columns))
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE3EA")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in df.itertuples(index=False):
        ws.append(list(row))
    ws.freeze_panes = "A5"
    _sheet_widths(ws, widths)
    wb.save(path)
    postprocess.scrub(path, producer, stamp)
    postprocess.finish_office(path, created, stamp + "T16:05:00Z", producer)


# ------------------------------------------------------------------ the registers
def dictionary(F):
    return f"""# Data dictionary

## {F["council"]}
The council's 2027 export from the network database: one row per station, pollutant and hour.
- site_code, site_name: the council station.
- pollutant: PM10, NO2, or NOx as NO2.
- date: dd/mm/yyyy (GMT); time_end: the end of the hour (01:00 to 24:00, the hour ending 24:00 closing the day).
- value: the hourly mean in ug/m3, blank where the hour is not valid.
- unit: ug/m3.
- status: R ratified and valid; F analyser fault; X rejected at ratification; S service or calibration; P power or
  telemetry lost; O station out of service (procedure section 2.3).

## {F["operator"]}
The terminal operator's 2027 PM10 workbook, one sheet per quarter of its logger's clock: a date/time column in UK clock
time (the start of the hour), then for each of its five stations a PM10 column (ug/m3) and a status column as ratified
under the procedure (OK valid; FAULT analyser fault; REJECTED rejected at ratification; SERVICE service or calibration;
POWER power or telemetry lost; OFF out of service).

## {F["met"]}
The port's met mast (10 m mast on the east breakwater): datetime_end_gmt (the end of the hour, GMT); wind_dir_deg (the
direction the wind blows from, degrees true, blank in calm); wind_speed_ms (m/s); rain_mm; air_temperature_c (degrees C). Hours the
mast did not report are blank.

## {F["register"]}
One row per station: site_code, site_name, operator, site_type, easting and northing (Ordnance Survey national grid,
metres), pm10_analyser, no2_analyser, inlet_height_m.

## {F["servicelog"]}
Both technicians' records: date and time (UK clock), site, site name, type (alarm, fault visit, service, calibration,
repair, site works, logger, deposit gauge), engineer, notes.

## {F["deposit"]}
The terminal's laboratory results for its deposit gauges: gauge, site, date on and date off (dd/mm/yyyy), days
exposed, insoluble, soluble and total deposit (mg), collecting area (m2) and deposition rate (mg/m2/day) as the
laboratory reports them, laboratory reference, remarks.

## {F["asr2026"]}
Table A.3 of last year's ASR: each council NO2 station's 2026 data capture and annual mean.

## {F["bgmap"]}
The background map extract used for the 2026 ASR: modelled annual mean background PM10 and NO2 for 2027 on the
national grid's 1 km squares (x and y, the square's centre, metres), with the main source sectors.
"""


def readme(F):
    return f"""# Thornmouth air quality network: 2027 data audit

Files for the PM10 data-capture review of the network's 2027 results, gathered by the council's air quality team.

- {F["request"]}: the air quality officer's request.
- {F["procedure"]}: the council's QA/QC procedure for the network (version 4).
- {F["council"]}: the council's 2027 export from the network database (ratified hourly PM10, NO2 and NOx).
- {F["operator"]}: the terminal operator's 2027 PM10 workbook.
- {F["met"]}: the port's met mast, 2027.
- {F["register"]}: the network's station register.
- {F["servicelog"]}: the council's and the operator's service records for 2027.
- {F["deposit"]}: the terminal's deposit gauge results.
- {F["plan"]}: the terminal's dust management and monitoring plan, revision 3.
- {F["asr2026"]}: last year's ASR NO2 table.
- {F["bgmap"]}: the background map extract used for last year's ASR.
- DATA_DICTIONARY.md: the fields of each file.
"""


def write_all(target, F, R):
    p = lambda k: os.path.join(target, F[k])
    request(p("request"), F)
    procedure(p("procedure"), F)
    plan(p("plan"), F)
    asr2026(p("asr2026"), F, R)
    operator_workbook(p("operator"), R)
    table_workbook(p("servicelog"), R.service_log(), "Service log 2027", "Network service log, 2027",
                   "Council stations: council technician. Terminal stations: the operator's technician. Times in UK clock time.",
                   COUNCIL, "2028-01-09", [11, 9, 7, 18, 13, 12, 110], "2027-01-04T08:30:00Z")
    table_workbook(p("deposit"), R.deposit_results(), "Results", "%s: deposit gauge results" % OPERATOR,
                   "Analysed by the terminal's contracted laboratory. Rates as reported by the laboratory.",
                   OPERATOR, "2028-01-19", [8, 13, 11, 11, 8, 10, 10, 12, 11, 13, 9, 34], "2027-01-12T10:10:00Z")
    with open(p("dict"), "w") as fh:
        fh.write(dictionary(F))
    with open(p("readme"), "w") as fh:
        fh.write(readme(F))
