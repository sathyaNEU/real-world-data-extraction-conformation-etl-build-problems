"""task245 goldens: the review workbook (each station's days above 50 over its valid days, its count for the whole year
with the range over the rebuild choices, the data-capture note, and the two audit figures beside it) and the one-page
note to the QA/QC auditor with its chart."""
import io

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess

XLSX = "thornmouth_pm10_2027_data_capture_review.xlsx"
PDF = "asr2027_pm10_data_capture_note.pdf"
ORG = "Thornmouth Borough Council"
AUTHOR = "Conor Robinson"


def _workbook(path, A):
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    ws["A1"] = "PM10 data-capture review of the 2027 results, with the NO2 and deposition figures for the audit"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Thornmouth air quality network, 2027. Prepared for the QA/QC audit of the 2027 data before the Annual "
                "Status Report; days and the note as defined in AQ-QP-04 v4, section 5.")
    ws["A2"].font = Font(italic=True, size=9)
    head = ["Site", "Site name", "Operator", "Valid daily means (of 365)", "Days above 50 ug/m3, valid daily means",
            "Days above 50 ug/m3, whole year", "Whole year, accepted range", "Data-capture note",
            "NO2 annual mean 2027 (ug/m3)", "Deposition rate 2027 (mg/m2/day)"]
    thin = Side(style="thin", color="A6ADB8")
    for j, h in enumerate(head, 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(bold=True, size=9.5)
        c.fill = PatternFill("solid", fgColor="DDE4EE")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = Border(bottom=thin)
    ws.row_dimensions[4].height = 58
    for i, r in enumerate(A["stations"], 5):
        vals = [r["code"], r["name"], r["operator"], r["valid_days"], r["valid_count"], r["year_count"],
                "%d to %d" % tuple(r["interval"]), "Note" if r["note"] else "",
                r["no2"] if r["no2"] is not None else None, r["dep"] if r["dep"] is not None else None]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(size=10, bold=(j == 8 and r["note"]))
            if j == 9:
                c.number_format = "0.0"
            elif j in (4, 5, 6, 10):
                c.number_format = "0"
            c.alignment = Alignment(horizontal="left" if j in (1, 2, 3) else ("center" if j in (7, 8) else "right"))
            if r["note"]:
                c.fill = PatternFill("solid", fgColor="FBEFD9")
    n = 5 + len(A["stations"])
    notes = [
        "Days above 50: daily means (00:00 to 24:00 GMT) of at least 18 valid hours. The whole-year count also counts the "
        "days without a valid daily mean (section 5.1): their hours are rebuilt from the median of the other stations' "
        "hourly means that hour plus the station's own mean increment over that median for the hour's wind (16-point "
        "sector, speed below or above 4 m/s, wet or dry day), taken from its valid hours; the day's valid hours are kept.",
        "Accepted range: the lowest to highest count with the other usual rebuild choices (4 to 36 sectors or a smooth "
        "fit in the wind direction, with or without speed and rain, the council's background stations, the rural station "
        "or a station-and-hour fit as the background), a day either side for day-boundary and thin-cell conventions. "
        "Every choice gives the same notes.",
        "NO2: valid hours of 2027, each hour once, at 20 C (1.91 ug/m3 per ppb; the loggers fitted in June 2027 "
        "convert at 1.88, so their hours are rescaled). Deposition: periods that started in 2027, the void period left "
        "out, each period's deposit over the collecting area of the gauge then fitted (the Frisbee gauges from the 7 June "
        "exchange), weighted by its days.",
        "Sources: network database export (ratified, 2027); Thornmouth Bulk Terminal PM10 workbook (6 March 2028); port "
        "met mast; station register; service log 2027; terminal deposit gauge results; terminal dust management and "
        "monitoring plan rev. 3."]
    for k, t in enumerate(notes):
        c = ws.cell(row=n + 1 + k, column=1, value=t)
        c.font = Font(size=8.5, italic=(k == len(notes) - 1))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=n + 1 + k, start_column=1, end_row=n + 1 + k, end_column=10)
        ws.row_dimensions[n + 1 + k].height = 40 if k < 3 else 28
    for col, w in zip("ABCDEFGHIJ", [7, 18, 22, 12, 14, 13, 14, 12, 13, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C5"
    ws.print_title_rows = "4:4"

    es = wb.create_sheet("Lost days")
    es["A1"] = "The noted stations' lost days and the wind"
    es["A1"].font = Font(bold=True, size=11)
    eh = ["Site", "Days without a valid daily mean", "of which the wind blew from the terminal for 6 hours or more",
          "Valid days with that wind: station mean (ug/m3)", "Valid days with that wind: network median (ug/m3)",
          "Lost days rebuilt above 50"]
    for j, h in enumerate(eh, 1):
        c = es.cell(row=3, column=j, value=h)
        c.font = Font(bold=True, size=9.5)
        c.fill = PatternFill("solid", fgColor="DDE4EE")
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    es.row_dimensions[3].height = 58
    for i, r in enumerate([r for r in A["stations"] if r["note"]], 4):
        vals = [r["code"], r["lost_days"], r["lost_bearing"], r["bearing_station"], r["bearing_network"], r["rebuilt_above"]]
        for j, v in enumerate(vals, 1):
            c = es.cell(row=i, column=j, value=v)
            c.number_format = "0" if j in (2, 3, 6) else "0.0"
    for col, w in zip("ABCDEF", [8, 16, 20, 18, 18, 14]):
        es.column_dimensions[col].width = w

    ns = wb.create_sheet("Notes")
    lines = [
        ("As of", "17 March 2028 (ratified 2027 data; the terminal's workbook of 6 March 2028)"),
        ("Year", "1 January to 31 December 2027, GMT days; hourly means stamped with the end of the hour"),
        ("Terminal data", "The operator's workbook is in UK clock time with each hour stamped at its start; moved to GMT "
                          "hour ending before any daily mean."),
        ("Note", "A station whose days above 50 among its valid daily means are 35 or fewer but whose whole-year count is "
                 "more than 35 (AQ-QP-04 v4, 5.2)."),
        ("Bearing", "The wind blowing from the terminal: a reported direction within 22.5 degrees of the bearing from the "
                    "station to the stockyard centre (the register's grid references; the plan's centre, E %d, N %d)."
                    % P.STOCKYARD_OS),
    ]
    for i, (k, v) in enumerate(lines, 1):
        ns.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = ns.cell(row=i, column=2, value=v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ns.row_dimensions[i].height = 30 if len(v) > 100 else 16
    ns.column_dimensions["A"].width = 14
    ns.column_dimensions["B"].width = 110
    wb.save(path)
    postprocess.scrub(path, AUTHOR, "2028-03-28")
    postprocess.finish_office(path, "2028-03-20T09:10:00Z", "2028-03-28T15:40:00Z", ORG)


def chart_png(A):
    rows = A["stations"]
    names = ["%s %s" % (r["code"], r["name"]) for r in rows]
    fig, ax = plt.subplots(figsize=(6.8, 4.3), dpi=200)
    y = list(range(len(rows)))[::-1]
    for yy, r in zip(y, rows):
        col = "#B5532B" if r["note"] else "#7A8CA5"
        ax.plot([r["valid_count"], r["year_count"]], [yy, yy], color=col, lw=1.6, zorder=1)
        ax.scatter([r["valid_count"]], [yy], s=22, facecolor="white", edgecolor=col, lw=1.2, zorder=2)
        ax.scatter([r["year_count"]], [yy], s=24, color=col, zorder=3)
        if r["note"]:
            ax.text(max(r["year_count"], r["interval"][1]) + 1.2, yy, "Note", va="center", fontsize=7, color="#B5532B",
                    fontweight="bold")
    ax.axvline(35, color="#1F2A3A", lw=1.1, ls="--")
    ax.text(35.6, len(rows) - 0.35, "35 days allowed", fontsize=7.2, color="#1F2A3A", va="center")
    ax.set_yticks(y)
    ax.set_yticklabels(names, fontsize=7.6)
    ax.set_xlim(0, max(r["year_count"] for r in rows) + 9)
    ax.set_ylim(-0.8, len(rows) - 0.1)
    ax.set_xlabel("Days in 2027 with a daily mean PM10 above 50 ug/m3 (open: valid daily means only; filled: whole year)",
                  fontsize=7.4)
    ax.set_title(A["chart_title"], fontsize=9.3, loc="left", fontweight="bold")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.tick_params(axis="x", labelsize=7.5)
    fig.text(0.01, 0.012, "Whole year: lost days rebuilt from the network's median and each station's own response to the "
                          "wind (AQ-QP-04 v4, 5.1 and 5.2).\nSource: network database (ratified 2027), terminal PM10 workbook, "
                          "port met mast.", fontsize=5.8, color="#555555", va="bottom")
    fig.tight_layout(rect=(0, 0.06, 1, 1))
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def _note(path, A):
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=9.2, leading=12.6, spaceAfter=5)
    head = ParagraphStyle("h", parent=body, fontName="Helvetica-Bold", fontSize=12.5, leading=15.5, spaceAfter=3)
    small = ParagraphStyle("s", parent=body, fontSize=7.6, leading=9.8, textColor=colors.HexColor("#444444"))
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=14 * mm,
                            bottomMargin=14 * mm, title=A["title"], author=AUTHOR, creator=ORG,
                            subject="PM10 data-capture review, 2027", invariant=1)
    s = [Paragraph("%s, Environmental Protection. To: Marc Bevan, Ashcroft Air Quality Services (QA/QC audit). "
                   "From: Conor Robinson. 28 March 2028." % ORG, small),
         Paragraph(A["title"], head)]
    for para in A["paras"]:
        s.append(Paragraph(para, body))
    img = Image(io.BytesIO(chart_png(A)), width=170 * mm, height=107 * mm)
    s += [img, Spacer(1, 2), Paragraph(A["footnote"], small)]
    doc.build(s)
    postprocess.scrub(path, ORG, "2028-03-28")


def write(folder, A):
    import os
    os.makedirs(folder, exist_ok=True)
    _workbook(os.path.join(folder, XLSX), A)
    _note(os.path.join(folder, PDF), A)
