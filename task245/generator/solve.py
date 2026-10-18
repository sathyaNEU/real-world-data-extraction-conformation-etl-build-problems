"""task245 answers from the shipped pack: the council's export and the operator's workbook brought onto one GMT
hour-ending grid, the met mast on the same grid, every reading of every station's whole-year count (estim.readings),
the call, and both asks. The goldens, figures.json and the write-up take their figures from here."""
import datetime as dt
import math
import os
import re

import numpy as np
import pandas as pd
from openpyxl import load_workbook

import estim as E
import params as P

H = P.HOURS
T0 = P.START
BST_ON = dt.datetime(2027, 3, 28, 1, 0)
BST_OFF = dt.datetime(2027, 10, 31, 1, 0)


def hour_index(t):
    """index of the hour that ENDS at GMT time t"""
    return int((t - T0).total_seconds() // 3600) - 1


def council(path):
    """PM10 (value, valid) for the council's stations, and the NO2 rows"""
    df = pd.read_csv(path, dtype=str, keep_default_na=False)
    df["d"] = pd.to_datetime(df["date"], format="%d/%m/%Y")
    hr = df["time_end"].str.slice(0, 2).astype(int)
    df["t"] = df["d"] + pd.to_timedelta(hr, unit="h")
    df["h"] = ((df["t"] - pd.Timestamp(T0)) / pd.Timedelta(hours=1)).astype(int) - 1
    return df


def operator(path):
    """the operator's quarterly sheets: UK clock time, hour beginning, per station value and status"""
    wb = load_workbook(path, data_only=True)
    rows = []
    seen_back = set()
    for ws in wb.worksheets:
        head = None
        for r in ws.iter_rows(values_only=True):
            if r and r[0] == "Date/time":
                head = list(r)
                continue
            if head is None or r[0] is None:
                continue
            t = dt.datetime.strptime(r[0], "%d/%m/%Y %H:%M")
            # the clock's hour beginning: BST from 28 March 02:00 local to 31 October 02:00 local; the repeated 01:00
            # on 31 October is BST the first time and GMT the second
            if t.date() == dt.date(2027, 10, 31) and t.hour == 1:
                key = t
                gmt = t - dt.timedelta(hours=1) if key not in seen_back else t
                seen_back.add(key)
            else:
                gmt = t - dt.timedelta(hours=1) if (dt.datetime(2027, 3, 28, 2) <= t < dt.datetime(2027, 10, 31, 1)) else t
            h = int((gmt - T0).total_seconds() // 3600)
            for k in range(1, len(head), 2):
                name = head[k].replace(" PM10", "")
                rows.append((name, h, r[k], r[k + 1]))
    return pd.DataFrame(rows, columns=["name", "h", "value", "status"])


def grid(target, F):
    """12 x 8760 PM10 values and validity, in the register's station order"""
    reg = pd.read_csv(os.path.join(target, F["register"]), dtype=str)
    codes = list(reg["site_code"])
    names = dict(zip(reg["site_name"], reg["site_code"]))
    xv = np.zeros((len(codes), H))
    valid = np.zeros((len(codes), H), bool)
    c = council(os.path.join(target, F["council"]))
    pm = c[c["pollutant"] == "PM10"]
    for code, g in pm.groupby("site_code"):
        i = codes.index(code)
        ok = g["status"] == "R"
        xv[i, g["h"][ok].values] = g["value"][ok].astype(float).values
        valid[i, g["h"][ok].values] = True
    o = operator(os.path.join(target, F["operator"]))
    for name, g in o.groupby("name"):
        i = codes.index(names[name])
        ok = g["status"] == "OK"
        xv[i, g["h"][ok].values] = g["value"][ok].astype(float).values
        valid[i, g["h"][ok].values] = True
    return codes, xv, valid


def met(target, F, calm="nearest", wet_mm=0.2):
    """the mast on the grid: direction (calm and missing hours given the nearest reported direction; calm 'vector':
    interpolated through the wind's unit vector instead), speed (missing hours interpolated), and the wet day (the
    mast's reported rain for the day totalling more than wet_mm)"""
    m = pd.read_csv(os.path.join(target, F["met"]), dtype=str, keep_default_na=False)
    t = pd.to_datetime(m["datetime_end_gmt"].str.replace("Z", ""), format="%Y-%m-%dT%H:%M")
    h = (((t - pd.Timestamp(T0)) / pd.Timedelta(hours=1)).astype(int) - 1).values
    theta = np.full(H, np.nan)
    speed = np.full(H, np.nan)
    rain = np.zeros(H)
    theta[h] = pd.to_numeric(m["wind_dir_deg"], errors="coerce").values
    speed[h] = pd.to_numeric(m["wind_speed_ms"], errors="coerce").values
    rain[h] = pd.to_numeric(m["rain_mm"], errors="coerce").fillna(0.0).values
    idx = np.arange(H)
    ok = ~np.isnan(theta)
    if calm == "vector":
        u = np.interp(idx, idx[ok], np.sin(np.radians(theta[ok])))
        v = np.interp(idx, idx[ok], np.cos(np.radians(theta[ok])))
        theta = np.degrees(np.arctan2(u, v)) % 360
    else:
        near = idx[ok][np.clip(np.searchsorted(idx[ok], idx), 0, ok.sum() - 1)]
        prev = idx[ok][np.clip(np.searchsorted(idx[ok], idx) - 1, 0, ok.sum() - 1)]
        pick = np.where(np.abs(prev - idx) <= np.abs(near - idx), prev, near)
        theta = theta[pick]
    oks = ~np.isnan(speed)
    speed[~oks] = np.interp(idx[~oks], idx[oks], speed[oks])
    wetday = rain.reshape(P.DAYS, 24).sum(1) > wet_mm
    return dict(theta=theta % 360, speed=speed, wet_h=np.repeat(wetday, 24))


def service_log(target, F):
    wb = load_workbook(os.path.join(target, F["servicelog"]), data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    head = list(rows[3])
    df = pd.DataFrame([r for r in rows[4:] if r[0]], columns=head)
    t = pd.to_datetime(df["Date"] + " " + df["Time (UK clock)"], format="%d/%m/%Y %H:%M")
    # UK clock to GMT: BST from 01:00 GMT on 28 March to 01:00 GMT on 31 October
    bst = (t >= pd.Timestamp("2027-03-28 02:00")) & (t < pd.Timestamp("2027-10-31 02:00"))
    df["gmt"] = t - pd.to_timedelta(bst.astype(int), unit="h")
    return df


def logger_fits(target, F):
    """each council station's new logger and its conversion, from the service log: {code: (first GMT hour, factor)}"""
    log = service_log(target, F)
    out = {}
    for _, r in log[log["Type"] == "Logger"].iterrows():
        fac = float(re.search(r"at ([0-9.]+) ug/m3 per ppb", r["Notes"]).group(1))
        out[r["Site"]] = (int((r["gmt"] - pd.Timestamp(T0)) / pd.Timedelta(hours=1)), fac)
    return out


def no2_means(target, F, resent="drop", fits="log"):
    """ask: each council NO2 station's 2027 annual mean at 20 C: valid hours, each hour once, and the hours its new
    logger converted at its own factor rescaled to the procedure's 1.91 ug/m3 per ppb (resent: 'drop' or 'keep';
    fits: 'log' from the service log, 'all' every hour rescaled, 'none' the export as it stands)"""
    c = council(os.path.join(target, F["council"]))
    n = c[(c["pollutant"] == "NO2") & (c["status"] == "R")]
    if resent == "drop":
        n = n.drop_duplicates(subset=["site_code", "h"])
    lf = logger_fits(target, F)
    out = {}
    for code, g in n.groupby("site_code"):
        v = g["value"].astype(float).values
        if fits != "none":
            h0, fac = lf[code]
            after = g["h"].values >= h0 if fits == "log" else np.ones(len(v), bool)
            v = np.where(after, v * P.NO2_PER_PPB / fac, v)
        out[code] = float(v.mean())
    return out


def frisbee_mm(target, F):
    """the plan's deposit gauge (Annex A): its collecting diameter in mm"""
    from pypdf import PdfReader
    txt = " ".join(" ".join(pg.extract_text().split()) for pg in PdfReader(os.path.join(target, F["plan"])).pages)
    return int(re.search(r"Frisbee-type deposit gauge, (\d+) mm collecting diameter", txt).group(1))


def frisbee_area(target, F):
    """the plan's deposit gauge as a collecting area in m2"""
    return math.pi * (frisbee_mm(target, F) / 2000.0) ** 2


def lab_area(target, F):
    """the collecting area the laboratory reports (one value throughout)"""
    wb = load_workbook(os.path.join(target, F["deposit"]), data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    col = list(rows[3]).index("Collecting area (m2)")
    vals = {float(r[col]) for r in rows[4:] if r[0]}
    assert len(vals) == 1, vals
    return vals.pop()


def gauge_swaps(target, F):
    """the exchange at which each terminal gauge became a Frisbee gauge, from the service log: {site code: date}"""
    log = service_log(target, F)
    sw = log[(log["Type"] == "Deposit gauge") & log["Notes"].str.contains("Frisbee")]
    return {r["Site"]: r["gmt"].date() for _, r in sw.iterrows()}


def deposition(target, F, area="log", window="started", void="out"):
    """ask: each terminal gauge's 2027 rate: periods started in 2027 (window 'ended': periods ended in 2027), void
    periods left out (void 'days': their days kept in the weights), each period's total deposit over the collecting area of
    the gauge exposed then and its days, weighted by days (area 'log' from the service log and the plan, 'lab' the
    laboratory's column, 'frisbee' the plan's gauge for every period); returns rates by site name"""
    wb = load_workbook(os.path.join(target, F["deposit"]), data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    head = rows[3]
    data = [dict(zip(head, r)) for r in rows[4:] if r[0]]
    reg = pd.read_csv(os.path.join(target, F["register"]), dtype=str)
    code_of = dict(zip(reg["site_name"], reg["site_code"]))
    fa = frisbee_area(target, F)
    sw = gauge_swaps(target, F)
    out = {}
    for site in sorted({d["Site"] for d in data}):
        num = den = 0.0
        for d in data:
            if d["Site"] != site:
                continue
            on = dt.datetime.strptime(d["Date on"], "%d/%m/%Y").date()
            off = dt.datetime.strptime(d["Date off"], "%d/%m/%Y").date()
            if (on.year if window == "started" else off.year) != 2027:
                continue
            days = int(d["Days exposed"])
            if (d["Remarks"] or "").startswith("VOID"):
                if void == "days":
                    den += days
                continue
            total = float(d["Total deposit (mg)"])
            lab_area = float(d["Collecting area (m2)"])
            if area == "lab":
                a = lab_area
            elif area == "frisbee":
                a = fa
            else:
                a = fa if on >= sw[code_of[site]] else lab_area
            num += total / a
            den += days
        out[site] = num / den
    return out


def lab_rates(target, F):
    """the laboratory's own rates, days-weighted over the periods that started in 2027, void periods out"""
    wb = load_workbook(os.path.join(target, F["deposit"]), data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    head = rows[3]
    data = [dict(zip(head, r)) for r in rows[4:] if r[0]]
    out = {}
    for site in sorted({d["Site"] for d in data}):
        num = den = 0.0
        for d in data:
            on = dt.datetime.strptime(d["Date on"], "%d/%m/%Y").date()
            if d["Site"] != site or on.year != 2027 or not d["Deposition rate (mg/m2/day)"]:
                continue
            num += float(d["Deposition rate (mg/m2/day)"]) * int(d["Days exposed"])
            den += int(d["Days exposed"])
        out[site] = num / den
    return out


def bearings(target, F):
    """each station's bearing to the stockyard centre (degrees, the direction a wind blows from to carry the stockyard
    to the station), from the register's grid references and the plan's centre"""
    from pypdf import PdfReader
    txt = " ".join(" ".join(pg.extract_text().split()) for pg in PdfReader(os.path.join(target, F["plan"])).pages)
    m = re.search(r"centred on national grid reference E (\d{6}), N (\d{6})", txt)
    e0, n0 = float(m.group(1)), float(m.group(2))
    reg = pd.read_csv(os.path.join(target, F["register"]), dtype=str)
    return {c: float(np.degrees(np.arctan2(e0 - float(e), n0 - float(n))) % 360)
            for c, e, n in zip(reg["site_code"], reg["easting"], reg["northing"])}


def reported_directions(target, F):
    """the mast's wind direction as reported, hour by hour (nan for a calm or missing hour)"""
    m = pd.read_csv(os.path.join(target, F["met"]), dtype=str, keep_default_na=False)
    t = pd.to_datetime(m["datetime_end_gmt"].str.replace("Z", ""), format="%Y-%m-%dT%H:%M")
    h = (((t - pd.Timestamp(T0)) / pd.Timedelta(hours=1)).astype(int) - 1).values
    theta = np.full(H, np.nan)
    theta[h] = pd.to_numeric(m["wind_dir_deg"], errors="coerce").values
    return theta


def answer(target, F):
    codes, xv, valid = grid(target, F)
    assert codes == P.CODES, "register order"
    W = met(target, F)
    R = E.readings(xv, valid, W)
    return dict(codes=codes, xv=xv, valid=valid, W=W, R=R, theta_reported=reported_directions(target, F),
                bearing=bearings(target, F))
