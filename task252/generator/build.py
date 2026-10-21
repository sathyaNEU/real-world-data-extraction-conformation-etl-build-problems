"""Cut the task252 pack from the world."""
import csv
import datetime as dt
import os

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import docs
import params as P
import postprocess
import records as R
import sim

FILES = dict(
    readings="nh3_online_readings_2027Q4.csv",
    lims="lims_ammonia_2027Q4.csv",
    checks="analyser_verification_2027Q4.xlsx",
    register="online_analyser_register.xlsx",
    standard="WQ-DS-04_online_monitoring_data_standard.pdf",
    process="water_treatment_process_summary_2027.pdf",
    dictionary="water_services_data_dictionary.md",
    notice="Kowhai_Bay_new_site_codes_from_1_November.eml",
    output="plant_daily_output_2027Q4.csv",
    rainfall="catchment_rainfall_daily_2027Q4.csv",
)
PACK = list(FILES)
DISTRACTORS = ("output", "rainfall")
CONTEXT = ("notice",)


def _csv(path, header, rows):
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)


def _ws(wb, title, header, rows, widths, first=False, fmts=None, wrap=()):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCE6EC")
        c.alignment = Alignment(vertical="top", wrap_text=True)
    for r in rows:
        ws.append(list(r))
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for col, f in (fmts or {}).items():
        j = header.index(col) + 1
        for i in range(2, len(rows) + 2):
            ws.cell(row=i, column=j).number_format = f
    for col in wrap:
        j = header.index(col) + 1
        for i in range(2, len(rows) + 2):
            ws.cell(row=i, column=j).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    return ws


def register(path, W):
    wb = Workbook()
    hdr = ["Analyser tag", "Plant", "Analyser", "Method", "Range (mg/L as N)", "Port sequence", "Sequence since",
           "Port time (min)", "Cycle (min)", "Check port", "Commissioned", "Last annual service"]
    comm = {"ASH": "2024-08-14", "BRN": "2024-09-03", "CAL": "2023-02-21", "DUN": "2024-08-27", "ELS": "2024-10-09",
            "FEN": "2024-11-12", "GAR": "2023-03-07", "HOL": "2023-03-15"}
    serv = {"ASH": "2027-08-17", "BRN": "2027-07-06", "CAL": "2027-03-23", "DUN": "2027-07-20", "ELS": "2027-08-31",
            "FEN": "2027-09-14", "GAR": "2027-05-11", "HOL": "2027-06-08"}
    rows, hist = [], []
    for c in P.CODES:
        x = W[c]
        seq = " > ".join(P.point(c, s) for s in x["order"])
        since = serv[c] if c in ("GAR", "HOL") else comm[c]
        rows.append([P.tag(c), P.NAME[c], "Colorimetric ammonia analyser, three-stream sequencer",
                     "Salicylate (indophenol blue)", "0 to 5", seq, dt.date.fromisoformat(since), 10, 30,
                     "%s-CHK" % c, dt.date.fromisoformat(comm[c]), dt.date.fromisoformat(serv[c])])
        if c in ("GAR", "HOL"):
            old = " > ".join(P.point(c, s) for s in ("raw", "fin", "fil"))
            hist.append([P.tag(c), dt.date.fromisoformat(comm[c]), dt.date.fromisoformat(serv[c]) - dt.timedelta(days=1),
                         old, "As commissioned"])
            hist.append([P.tag(c), dt.date.fromisoformat(serv[c]), None, seq,
                         "Resequenced at annual service (filtered water read ahead of final water)"])
        else:
            hist.append([P.tag(c), dt.date.fromisoformat(comm[c]), None, seq, "As commissioned"])
    _ws(wb, "Analysers", hdr, rows, [16, 13, 44, 26, 15, 34, 14, 14, 11, 12, 14, 18], first=True,
        fmts={"Sequence since": "yyyy-mm-dd", "Commissioned": "yyyy-mm-dd", "Last annual service": "yyyy-mm-dd"})
    pts = []
    for c in P.CODES:
        for s in ("raw", "fil", "fin"):
            lp = R.LIMS_RAW[c] if s == "raw" else R.LIMS_NEW[(c, s)]
            pts.append([P.point(c, s), P.NAME[c], P.WATER[s], lp, R.LOCATION[s]])
        pts.append(["%s-CHK" % c, P.NAME[c], "Check standard", "", "Analyser check port (manual)"])
    _ws(wb, "Sample points", ["Sample point", "Plant", "Water", "Laboratory code", "Location"], pts, [14, 13, 15, 16, 44])
    _ws(wb, "Port sequence history", ["Analyser tag", "From", "To", "Port sequence", "Note"], hist, [16, 12, 12, 34, 60],
        fmts={"From": "yyyy-mm-dd", "To": "yyyy-mm-dd"})
    wb.save(path)


def checks(path, W):
    wb = Workbook()
    rows = []
    for r in R.checks(W):
        rows.append([r["tag"], r["plant"], r["start"].replace(tzinfo=None), P.CHECK_STD, *r["reads"],
                     round(r["mean"], 4), r["result"], r["by"]])
    hdr = ["Analyser tag", "Plant", "Check started (NZ time)", "Standard (mg/L as N)", "Reading 1", "Reading 2",
           "Reading 3", "Mean", "Within 10%", "Checked by"]
    _ws(wb, "Checks Q4 2027", hdr, rows, [16, 13, 22, 20, 11, 11, 11, 10, 11, 18], first=True,
        fmts={"Check started (NZ time)": "dd/mm/yyyy hh:mm", "Standard (mg/L as N)": "0.000", "Reading 1": "0.000",
              "Reading 2": "0.000", "Reading 3": "0.000", "Mean": "0.0000"})
    ws = wb.create_sheet("Procedure")
    for line in ["Monthly verification of the online ammonia analysers (WQ-DS-04 section 5).",
                 "Standard: 0.100 mg/L as N ammonium, made up fresh from the 1000 mg/L stock on the day.",
                 "Switch the sequencer to the check port, flush the cell with the standard, take three readings.",
                 "Pass: the mean of the three within 10 per cent of the standard (0.090 to 0.110).",
                 "Return the sequencer to its sample ports and confirm the next raw water reading.",
                 "After clearing any analyser fault, run the same check before returning the sequencer to its sample "
                 "ports (not recorded on this sheet)."]:
        ws.append([line])
    ws.column_dimensions["A"].width = 100
    wb.save(path)


def run(target):
    os.makedirs(target, exist_ok=True)
    W = sim.build()
    F = {k: os.path.join(target, v) for k, v in FILES.items()}
    _csv(F["readings"], ["site_code", "analyser_tag", "reading_end_utc", "sample_point", "nh3_n_mgl", "qc"],
         R.readings(W))
    L = R.lims(W)
    cols = ["sample_no", "lims_point", "location", "collected", "received", "analysed", "released", "determinand",
            "result", "units", "lod"]
    _csv(F["lims"], cols, [[r[c] for c in cols] for r in sorted(L, key=lambda r: r["sample_no"])])
    checks(F["checks"], W)
    register(F["register"], W)
    docs.standard(F["standard"])
    docs.process(F["process"], W)
    docs.dictionary(F["dictionary"])
    docs.notice(F["notice"], W)
    rng = np.random.default_rng(P.SEED_DOCS)
    _csv(F["output"], ["date", "site_code", "plant", "supplied_ML"], R.output(W, rng))
    _csv(F["rainfall"], ["date", "gauge", "gauge_name", "rain_mm"], R.rainfall(W, rng))
    postprocess.scrub(F["standard"], "Microsoft Word 2019", "2027-06-02T09:14:00")
    postprocess.scrub(F["process"], "Microsoft Word 2019", "2027-07-19T15:41:00")
    for k, made, saved in (("checks", "2027-10-04T20:31:00Z", "2028-01-06T22:12:00Z"),
                           ("register", "2023-02-14T03:05:00Z", "2027-11-02T01:48:00Z")):
        postprocess.scrub(F[k], P.ORG, saved[:10])
        postprocess.finish_office(F[k], made, saved, P.ORG)
    return W


if __name__ == "__main__":
    import sys
    run(sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                                                           "target"))
