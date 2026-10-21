"""task252 golden, computed from the pack alone: each treated channel's quarter offset (carryover from the two ports
before it, fitted apart from the raw water's lagged breakthrough, plus the analyser's calibration error at the
channel's level, its form read from the check standard against the raw grabs with the raw port's own carryover taken
out), the flags, and the two asks of the return. Every licensed and unlicensed route is a switch on one function."""
import csv
import datetime as dt
import os
import re

import numpy as np
from openpyxl import load_workbook

import params as P

LAGS = list(range(12, 49, 3))        # raw-water lags in slots: two to eight hours, the plants' transit times


def _t(s):
    return dt.datetime.strptime(s, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=P.UTC)


def load(target):
    """the readings as slot-indexed arrays per plant: the G value (a reloaded slot's G row stands), the slot's code
    as the export shows it, the stream each slot read, the check readings"""
    path = os.path.join(target, "nh3_online_readings_2027Q4.csv")
    D = {c: dict(val=np.full(P.NS, np.nan), v=np.full(P.NS, np.nan), f=np.zeros(P.NS, bool), g=np.zeros(P.NS, bool),
                 present=np.zeros(P.NS, bool), point=np.array([""] * P.NS, dtype=object)) for c in P.CODES}
    with open(path) as fh:
        for row in csv.DictReader(fh):
            x = D[row["site_code"]]
            i = int(round((_t(row["reading_end_utc"]) - P.T0).total_seconds() / 600)) - 1
            x["present"][i] = True
            x["point"][i] = row["sample_point"]
            if row["qc"] == "G":
                x["val"][i] = float(row["nh3_n_mgl"]); x["g"][i] = True
            elif row["qc"] == "V":
                x["v"][i] = float(row["nh3_n_mgl"])
            elif row["qc"] == "F":
                x["f"][i] = True
    reg = load_workbook(os.path.join(target, "online_analyser_register.xlsx"), read_only=True)
    for r in list(reg["Analysers"].iter_rows(values_only=True))[1:]:
        code = r[0].split("-")[1]
        D[code]["order"] = [{"RW1": "raw", "FI1": "fil", "FW1": "fin"}[p.strip().split("-")[1]] for p in r[5].split(">")]
    for c in P.CODES:
        x = D[c]
        x["sl"] = np.array([x["order"][i % 3] for i in range(P.NS)])
        cell = x["val"].copy()
        m = ~np.isnan(x["v"])
        cell[m] = x["v"][m]                         # the cell held the standard on a check cycle
        x["cell"] = cell
    return D


def _rawfull(x):
    idx = np.where((x["sl"] == "raw") & x["g"])[0]
    return np.interp(np.arange(P.NS), idx, x["val"][idx])


def fit(x, s, how="G", lags_slots=None, form="mix"):
    """carryover into stream s: 'G' the two slots before with the raw water at the transit lags (licensed); 'RAW' the
    raw reading of the same cycle with the lags (licensed for the channel straight after raw, and two slots back for
    the third port); 'P1' the slot before only; 'D1' first differences on the slot before"""
    idx = np.where((x["sl"] == s) & x["g"])[0]
    idx = idx[idx >= 48]
    p1 = x["cell"][idx - 1]; p2 = x["cell"][idx - 2]
    ok = ~np.isnan(p1) & ~np.isnan(p2)
    idx, p1, p2 = idx[ok], p1[ok], p2[ok]
    r = x["val"][idx]
    rf = _rawfull(x)
    lags = [rf[idx - L] for L in (lags_slots or LAGS)]
    one = np.ones(idx.size)
    if how == "G":
        c = np.linalg.lstsq(np.column_stack([one, p1, p2] + lags), r, rcond=None)[0]
        k1, k2 = c[1], c[2]
        co = k1 * np.mean(p1 - r) + k2 * np.mean(p2 - r)
        if form == "mix":
            co /= (1 - k1 - k2)
        return dict(k1=k1, k2=k2, co=co)
    if how == "RAW":
        back = x["order"].index(s)
        q = rf[idx - back]
        c = np.linalg.lstsq(np.column_stack([one, q] + lags), r, rcond=None)[0]
        return dict(k1=c[1], k2=0.0, co=c[1] * np.mean(q - r) / (1 - c[1]))
    if how == "P1":
        c = np.linalg.lstsq(np.column_stack([one, p1] + lags), r, rcond=None)[0]
        return dict(k1=c[1], k2=0.0, co=c[1] * np.mean(p1 - r) / (1 - c[1]))
    if how == "D1":
        dr = np.diff(r); dp = np.diff(p1)
        c = np.linalg.lstsq(np.column_stack([np.ones(dr.size), dp]), dr, rcond=None)[0]
        return dict(k1=c[1], k2=0.0, co=c[1] * np.mean(p1 - r))
    if how == "CRUDE":                   # the coefficient on the same cycle's raw reading times its mean
        back = x["order"].index(s)
        q = rf[idx - back]
        c = np.linalg.lstsq(np.column_stack([one, q] + lags), r, rcond=None)[0]
        return dict(k1=c[1], k2=0.0, co=c[1] * np.mean(q))
    if how == "FIRST":                   # carryover only into the port read straight after raw water
        if x["order"].index(s) != 1:
            return dict(k1=0.0, k2=0.0, co=0.0)
        return fit(x, s, "RAW")
    raise ValueError(how)


CARRY = "POOL"                         # the golden's carryover estimate


def pooled(x, lags_slots=None, form="mix"):
    """carryover as a mix in the analyser's cell: each treated reading is (1 - k1 - k2) of its own water, k1 of the
    slot before and k2 of the one before that. k1 and k2 belong to the analyser, so they are fitted once on both
    treated ports together, each port with its own level and its own response to the raw water at the transit lags;
    each channel's carryover offset is k1 and k2 of the differences between the two slots before it and its own
    water. form 'add' is the additive law, k times the readings before (a rival: the raw ports reading low against
    the raw grabs rule it out)"""
    rf = _rawfull(x)
    lg = lags_slots or LAGS
    nl = len(lg)
    Xs, Ys, keep = [], [], {}
    for j, s in enumerate(("fil", "fin")):
        idx = np.where((x["sl"] == s) & x["g"])[0]
        idx = idx[idx >= 48]
        p1 = x["cell"][idx - 1]; p2 = x["cell"][idx - 2]
        ok = ~np.isnan(p1) & ~np.isnan(p2)
        idx, p1, p2 = idx[ok], p1[ok], p2[ok]
        r = x["val"][idx]
        X = np.zeros((idx.size, 4 + 2 * nl))
        X[:, j] = 1.0
        X[:, 2] = p1; X[:, 3] = p2
        X[:, 4 + j * nl:4 + (j + 1) * nl] = np.column_stack([rf[idx - L] for L in lg])
        Xs.append(X); Ys.append(r)
        keep[s] = (p1, p2, r)
    c = np.linalg.lstsq(np.vstack(Xs), np.concatenate(Ys), rcond=None)[0]
    k1, k2 = float(c[2]), float(c[3])
    co = {}
    for s, (p1, p2, r) in keep.items():
        if form == "add":
            co[s] = k1 * float(np.mean(p1)) + k2 * float(np.mean(p2))
        else:
            co[s] = (k1 * float(np.mean(p1 - r)) + k2 * float(np.mean(p2 - r))) / (1 - k1 - k2)
    return dict(k1=k1, k2=k2, co=co)


def checks(target):
    wb = load_workbook(os.path.join(target, "analyser_verification_2027Q4.xlsx"), read_only=True)
    out = {c: [] for c in P.CODES}
    for r in list(wb["Checks Q4 2027"].iter_rows(values_only=True))[1:]:
        code = r[0].split("-")[1]
        out[code] += [float(v) for v in r[4:7]]
    return out


def lims(target):
    """the laboratory's rows with the plant and water read from the register's codes, the notice's old codes and
    the location text"""
    rows = []
    old = {}
    with open(os.path.join(target, "Kowhai_Bay_new_site_codes_from_1_November.eml")) as fh:
        for a, b in re.findall(r"^\s+(TW-\S+)\s+(TUR\S+)\s*$", fh.read(), flags=re.M):
            old[a] = b
    reg = load_workbook(os.path.join(target, "online_analyser_register.xlsx"), read_only=True)
    code_of = {}
    for r in list(reg["Sample points"].iter_rows(values_only=True))[1:]:
        if r[3]:
            pl = r[0].split("-")[0]
            code_of[r[3]] = (pl, {"RW1": "raw", "FI1": "fil", "FW1": "fin"}[r[0].split("-")[1]])
    for a, b in old.items():
        code_of[a] = code_of[b]
    with open(os.path.join(target, "lims_ammonia_2027Q4.csv")) as fh:
        for row in csv.DictReader(fh):
            pl, s = code_of[row["lims_point"]]
            f = lambda v: dt.datetime.strptime(v, "%d/%m/%Y %H:%M")
            res = row["result"]
            rows.append(dict(no=row["sample_no"], code=pl, stream=s, collected=f(row["collected"]),
                             received=f(row["received"]), analysed=f(row["analysed"]), released=f(row["released"]),
                             cens=res.startswith("<"), value=float(res.lstrip("<"))))
    return rows


def raw_pairs(x, L, code):
    """each raw grab against the raw port's G reading of the slot it was collected in, or the nearest one, with the
    readings of the two slots before that reading"""
    idx_raw = np.where((x["sl"] == "raw") & x["g"])[0]
    idx_raw = idx_raw[(idx_raw >= 2) & ~np.isnan(x["cell"][np.maximum(idx_raw - 1, 0)])
                      & ~np.isnan(x["cell"][np.maximum(idx_raw - 2, 0)])]
    out = []
    for g in L:
        if g["code"] != code or g["stream"] != "raw":
            continue
        t = g["collected"].replace(tzinfo=P.NZ)
        i = P.slot_of(t)
        j = idx_raw[np.argmin(np.abs(idx_raw - i))]
        out.append((g["value"], x["val"][j], x["cell"][j - 1], x["cell"][j - 2]))
    return np.array(out)


def attribute(x, L, code, chk, corrected=True, k=None):
    """the check's shortfall read as a zero or a span error: under a zero error the raw grabs show the same absolute
    shortfall as the check, under a span error one in proportion; corrected takes the raw port's own carryover from
    the two ports before it out first"""
    dev = float(np.mean(chk)) - P.CHECK_STD
    A = raw_pairs(x, L, code)
    lab, rd, p1, p2 = A.T
    d = rd - lab
    if corrected:
        k1, k2 = k
        d = d - (k1 * (p1 - lab) + k2 * (p2 - lab))
    so = float(np.sum((d - dev) ** 2)); sg = float(np.sum((d - dev / P.CHECK_STD * lab) ** 2))
    return ("zero" if so < sg else "span"), dev, so, sg


def offsets(target, carry="G", check="corr", D=None, L=None, C=None, **fit_kw):
    """check: 'corr' (the golden), 'naive' (the two-point read without the raw port's carryover), 'span' (every
    shortfall a span error), 'zero' (every shortfall added as an offset), 'none' (checks passed, nothing done)"""
    D = D or load(target); L = L or lims(target); C = C or checks(target)
    out, info = {}, {}
    for code in P.CODES:
        x = D[code]
        if carry == "POOL":
            pk = pooled(x, **fit_kw)
            fits = {s: dict(k1=pk["k1"], k2=pk["k2"], co=pk["co"][s]) for s in ("fil", "fin")}
        else:
            fits = {s: fit(x, s, carry, **fit_kw) for s in ("fil", "fin")}
        pk = pooled(x)
        k = (pk["k1"], pk["k2"])
        dev = float(np.mean(C[code])) - P.CHECK_STD
        kind = None
        if abs(dev) >= 0.003 and check in ("corr", "naive"):
            kind = attribute(x, L, code, C[code], corrected=(check == "corr"), k=k)[0]
        elif abs(dev) >= 0.003:
            kind = {"span": "span", "zero": "zero", "none": None}[check]
        info[code] = dict(k=k, dev=dev, kind=kind)
        for s in ("fil", "fin"):
            idx = np.where((x["sl"] == s) & x["g"])[0]
            level = float(np.mean(x["val"][idx]))
            cal = 0.0 if kind is None else (dev if kind == "zero" else dev / P.CHECK_STD * level)
            out[(code, s)] = fits[s]["co"] + cal
    return out, info


def flags(off):
    return sorted(k for k, v in off.items() if abs(v) > P.THRESH)


def name(k):
    return P.point(k[0], k[1])


# ---- ask A: samples analysed more than 24 hours after collection --------------------------------------------------

def clock_bracket(L):
    """the minutes every analysis stamp must move forward to sit after reception logged its bottle and before the
    laboratory released its result: (the largest reception less stamp, the smallest release less stamp)"""
    lo = max((g["received"] - g["analysed"]).total_seconds() / 60 for g in L)
    hi = min((g["released"] - g["analysed"]).total_seconds() / 60 for g in L)
    return lo, hi


def ask_a(target, L=None, shift_h=1.0, scope=("fil", "fin"), current_codes_only=False, dedup="bottle", fix_early=False):
    """per plant, the filtered and final water samples analysed more than 24 hours after collection. The analysis
    stamps run behind the laboratory's other clocks: runs are stamped before reception logged their bottles, none is
    released within an hour of its stamp, and every shift inside that bracket counts the same, so the stamps move
    forward an hour (shift_h=1: the analyser on standard time, the rest on daylight time). A sample is a bottle: rows
    with the same point, analysis stamp and result are one bottle logged twice ('bottle'); 'none' keeps every row,
    'point_time' one row per point and collection time, 'content' drops only rows equal in every field but the sample
    number. fix_early moves only the stamps that sit before reception, to reception"""
    L = L or lims(target)
    out = {c: 0 for c in P.CODES}
    seen = set()
    for g in L:
        if g["stream"] not in scope:
            continue
        if current_codes_only and g["collected"] < dt.datetime(2027, 11, 1):
            continue
        key = {"bottle": (g["code"], g["stream"], g["analysed"], g["value"], g["cens"]),
               "point_time": (g["code"], g["stream"], g["collected"]),
               "content": (g["code"], g["stream"], g["collected"], g["received"], g["analysed"], g["value"], g["cens"]),
               "none": g["no"]}[dedup]
        if key in seen:
            continue
        seen.add(key)
        t = g["analysed"] + dt.timedelta(hours=shift_h)
        if fix_early:
            t = max(g["analysed"], g["received"])
        h = (t - g["collected"]).total_seconds() / 3600
        if h > P.HOLD_H:
            out[g["code"]] += 1
    return out


# ---- ask B: hours at the fixed dose -------------------------------------------------------------------------------

def runs(x):
    """runs of F rows with nothing filled beneath them, chained across gaps of up to two slots so one port's run, every
    third slot, stays one run"""
    bad = np.where(x["f"] & ~x["g"])[0]
    if bad.size == 0:
        return []
    cuts = np.where(np.diff(bad) > 3)[0]
    return [(int(a), int(b)) for a, b in zip(np.r_[bad[0], bad[cuts + 1]], np.r_[bad[cuts], bad[-1]])]


def checked(x, e):
    """a check-port reading in the three slots after slot e"""
    return bool(np.any(~np.isnan(x["v"][e + 1:min(e + 4, P.NS)])))


def whole_plant(x, a, b):
    """every slot from a to b an F row with nothing filled: the plant's whole record, every port"""
    return bool(np.all(x["f"][a:b + 1] & ~x["g"][a:b + 1]))


def telemetry(x, rule="plc"):
    """the slots of the unfilled runs that were the historian's alone, the PLC reading its analyser throughout. 'plc'
    (the golden): every port unfilled and no check after it, since a failed poll takes the plant's whole record and a
    fault is followed by the technician's check, while a run at one port with the others read either side is a reading
    the PLC never had; 'checked': any run with no check after it; 'allport': any run at every port"""
    out = np.zeros(P.NS, bool)
    for a, b in runs(x):
        hit = {"plc": whole_plant(x, a, b) and not checked(x, b), "checked": not checked(x, b),
               "allport": whole_plant(x, a, b)}[rule]
        if hit:
            out[a:b + 1] = True
    return out


def fil_good(x, how):
    """the filtered water slots with a reading the PLC had, as each path reads the export"""
    fi = np.where(x["sl"] == "fil")[0]
    if how in ("plc", "nohold", "checked", "allport"):
        tel = telemetry(x, "plc" if how == "nohold" else how)
        good = (~(x["f"][fi] & ~x["g"][fi]) | tel[fi]) & np.isnan(x["v"][fi])
    elif how == "faults":
        good = ~(x["f"][fi] & ~x["g"][fi]) & np.isnan(x["v"][fi])
    elif how == "historian":
        good = x["g"][fi]
    elif how == "allF":
        good = ~x["f"][fi] & np.isnan(x["v"][fi])
    else:
        raise ValueError(how)
    return fi[good]


def ask_b(target, D=None, how="plc"):
    """per plant, hours the chlorine dose ran at its fixed dose: the PLC holds the last good filtered reading for two
    hours, then doses fixed until the next good filtered reading. 'plc' (the golden): the readings the PLC had, a
    failed poll's filled G row, the historian's lost hours and the whole-plant runs nothing filled and no check
    followed (an outstation change) all good at the plant; 'checked': every unfilled run with no check after it good
    (a filtered port's lost signal read as the telemetry's); 'allport': every whole-plant run good (a reagent or cell
    fault read as the telemetry's); 'faults': every F row with nothing filled a gap; 'historian': gaps in the
    historian's G rows; 'allF': every F row a gap, the lost hours not; 'nohold': the plc's gaps summed whole"""
    D = D or load(target)
    out = {}
    for code in P.CODES:
        x = D[code]
        gaps = np.diff(fil_good(x, how)) * P.SLOT_MIN / 60.0
        if how == "nohold":
            hrs = float(np.sum(gaps[gaps > 0.5] - 0.5))
        else:
            hrs = float(np.sum(np.clip(gaps - P.DOSE_HOLD_H, 0, None)))
        out[code] = hrs
    return out


def golden(target):
    D = load(target); L = lims(target); C = checks(target)
    off, info = offsets(target, CARRY, "corr", D, L, C)
    return dict(off=off, info=info, flags=flags(off), ask_a=ask_a(target, L), ask_b=ask_b(target, D))


if __name__ == "__main__":
    import sys
    tg = sys.argv[1]
    G = golden(tg)
    for k, v in sorted(G["off"].items()):
        print("%-8s %+.4f %s" % (name(k), v, "FLAG" if k in G["flags"] else ""))
    for c, i in G["info"].items():
        print(c, "k1 %.3f k2 %.3f dev %+.4f %s" % (i["k"][0], i["k"][1], i["dev"], i["kind"]))
    print("ask A", G["ask_a"])
    print("ask B", {c: round(v, 1) for c, v in G["ask_b"].items()})
