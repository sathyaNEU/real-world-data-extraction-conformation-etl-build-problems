"""Render the two golden deliverables from the figures recomputed off target/."""
import datetime as dt
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import golden as G

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "golden"
D = dt.date
ORDER = ["Parts & van stock", "Engineer attendance & punctuality", "Travel & routing", "Job duration & standards",
         "Customer access", "Other", "Rota & cover"]
SHORT = {"Parts & van stock": "Parts", "Engineer attendance & punctuality": "Punctuality", "Travel & routing": "Travel",
         "Job duration & standards": "Job duration", "Customer access": "Access", "Other": "Other",
         "Rota & cover": "Rota & cover"}
COLOURS = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#4a3aa7"]
PATCHES = ["Glasgow North & West", "Glasgow South & East", "Lanarkshire", "Ayrshire & Renfrewshire",
           "Edinburgh & Lothians", "Fife & Forth Valley"]
SURFACE = "#fcfcfb"


def chart(fig, path):
    weeks = [k for k in sorted(fig["weekly"]) if "2026-01-01" <= k <= "2026-09-24"]
    x = list(range(len(weeks)))
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10})
    f, ax = plt.subplots(figsize=(13, 6.6), dpi=150)
    f.patch.set_facecolor(SURFACE)
    ax.set_facecolor(SURFACE)
    bottom = [0] * len(weeks)
    for fam, col in zip(ORDER, COLOURS):
        vals = [fig["weekly"][w].get(fam, 0) for w in weeks]
        ax.bar(x, vals, bottom=bottom, width=0.72, color=col, edgecolor=SURFACE, linewidth=1.2, label=fam,
               zorder=3)
        bottom = [b + v for b, v in zip(bottom, vals)]
    tot = bottom
    # go-live marker: 1 July is the Wednesday that closes the week of 25 June
    i_jul = weeks.index("2026-07-02")
    xm = i_jul - 0.5
    ax.axvline(xm, color="#52514e", linewidth=1.2, zorder=4)
    top = max(tot) * 1.16
    ax.text(xm - 0.3, top * 0.965, "Pathline 4.2 go-live, 1 Jul", ha="right", va="top", fontsize=9.5,
            color="#0b0b0b")
    # worst week
    wi = max(range(len(weeks)), key=lambda i: tot[i])
    wd = D.fromisoformat(weeks[wi])
    ax.annotate("Week of %d %s: %d missed" % (wd.day, wd.strftime("%b"), tot[wi]), xy=(wi, tot[wi] + 0.8),
                xytext=(wi, tot[wi] + 9), fontsize=9.5, color="#0b0b0b", ha="center", va="bottom",
                arrowprops=dict(arrowstyle="-", color="#52514e", linewidth=1))
    labs = []
    for i, w in enumerate(weeks):
        d = D.fromisoformat(w)
        labs.append("%d %s" % (d.day, d.strftime("%b")) if i % 4 == 0 else "")
    ax.set_xticks(x)
    ax.set_xticklabels(labs, fontsize=9, color="#52514e")
    ax.tick_params(axis="x", length=0)
    ax.tick_params(axis="y", colors="#52514e", length=0)
    ax.set_ylim(0, top)
    ax.set_xlim(-0.7, len(weeks) - 0.3)
    ax.grid(axis="y", color="#e4e3df", linewidth=0.8, zorder=0)
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.spines["bottom"].set_color("#c9c8c2")
    ax.set_ylabel("Missed appointments", color="#52514e")
    ax.set_xlabel("Reporting week (Thursday to Wednesday), labelled by its first day", color="#52514e")
    f.text(0.012, 0.965, "Late out-of-hours call-outs drove the summer rise and stop with the winter panel; "
           "van stock is what carries into January", ha="left", va="top", fontsize=13, fontweight="bold",
           color="#0b0b0b")
    f.text(0.012, 0.918, "Scotland region, missed appointments each week by the family they trace back to, "
           "January to September 2026", ha="left", va="top", fontsize=10, color="#52514e")
    handles, labels = ax.get_legend_handles_labels()
    leg = f.legend(handles, labels, ncol=7, loc="upper left", bbox_to_anchor=(0.008, 0.885), frameon=False,
                   fontsize=9, handlelength=1.1, columnspacing=1.2)
    for t in leg.get_texts():
        t.set_color("#0b0b0b")
    f.tight_layout(rect=(0, 0, 1, 0.84))
    f.savefig(path, facecolor=SURFACE, metadata={"Software": None})
    plt.close(f)
    return weeks[wi], tot[wi]


def _shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcol)
    tcPr.append(shd)


def _table(doc, header, rows, widths, bold_last=False):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, header):
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9)
        _shade(c, "E8E8E8")
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, (c, v) in enumerate(zip(cells, row)):
            c.text = ""
            p = c.paragraphs[0]
            run = p.add_run(str(v))
            run.font.size = Pt(9)
            if bold_last and i == len(rows) - 1:
                run.bold = True
            if j > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            c.width = Cm(w)
    return t


def _footer(doc, text):
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    p.text = ""
    r = p.add_run(text + "    Page ")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    run = p.add_run()
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)
    run.font.size = Pt(8)


def _note(doc, text):
    q = doc.add_paragraph()
    r = q.add_run(text)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    return q


def memo(fig, path, worst):
    fc = fig["family_counts"]
    cred = fig["credits_by_family"]
    comp = fig["complaints_by_family"]
    grid = fig["patch_family"]
    total = sum(fc.values())
    num = lambda n: "{:,}".format(n)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for sct in doc.sections:
        sct.left_margin = sct.right_margin = Cm(2.2)
        sct.top_margin = Cm(1.8)
        sct.bottom_margin = Cm(1.6)
    _footer(doc, "Kelvane Home Services, Field Operations Scotland. Internal.")
    h = doc.add_paragraph()
    r = h.add_run("Missed appointments into the peak: fix van stock before January")
    r.bold = True
    r.font.size = Pt(14)
    m = doc.add_paragraph()
    r = m.add_run("To: Ryan Webster, Resource Planning desk    From: Gillian Sanderson, Field Operations, Scotland\n"
                  "Date: 6 October 2026    Ref: FS-07 quarterly review, July to September 2026")
    r.font.size = Pt(9.5)
    p = doc.add_paragraph()
    r = p.add_run("Cause to fix before January: Parts & van stock. ")
    r.bold = True
    p.add_run("%d of the appointments we missed between 1 July and 30 September trace back to an engineer having to "
              "collect a part at a merchant on the way to a job, which has happened far more often since the vans "
              "went to fortnightly top-ups at the start of July. Of the causes that will still be with us over "
              "January to March it is the largest, and nothing about the winter touches it." % fc["Parts & van stock"])
    p = doc.add_paragraph()
    r = p.add_run("Change before January: ")
    r.bold = True
    p.add_run("put the vans back on weekly top-ups before the peak, starting with the diverter valves and PCBs that "
              "send engineers to the trade counter.")
    p = doc.add_paragraph()
    r = p.add_run("Why not Rota & cover, the biggest cause this quarter. ")
    r.bold = True
    sb = fig["standby_in_post"]
    p.add_run("%d of the quarter's misses were day-roster engineers who took an evening call-out that cleared after "
              "21:00 and could not log on until eleven hours later (HR-FS-03 3.1, with call-out time counted as "
              "working time under 6.3). That only happens when the out-of-hours desk sends evening calls to the "
              "call-out list, which it has done all summer with the standby pool down from %d to %d. From October "
              "to May the contractor panel takes the evening calls standby cannot reach: last winter no call-out "
              "list attendance cleared after 21:00 and there was not one Rota & cover miss between October and "
              "May. The panel has been attending again since 1 October. Standby still needs rebuilding before next "
              "summer, but it will not keep a January appointment."
              % (fc["Rota & cover"], sb["2025-07-01"], sb["2026-09-30"]))
    p = doc.add_paragraph()
    r = p.add_run("Second over the peak: Engineer attendance & punctuality, %d. " % fc["Engineer attendance & punctuality"])
    r.bold = True
    p.add_run("These are late log-ons with no call-out behind them, and they ran no higher last winter than "
              "last summer. Leslie is right that Pathline 4.2 made the long legs tighter and Travel & routing has "
              "grown since the go-live, but at %d it sits below both." % fc["Travel & routing"])
    p = doc.add_paragraph()
    r = p.add_run("By cause family, July to September 2026")
    r.bold = True
    rows = []
    for fam in sorted(ORDER, key=lambda f: -fc[f]):
        rows.append([fam, num(fc[fam]), num(cred.get(fam, 0)), comp.get(fam, 0)])
    _table(doc, ["Cause family", "Missed appointments", "Goodwill credits (GBP)", "Complaints"], rows,
           [6.0, 3.6, 4.2, 2.6])
    _note(doc, "Source: Pathline appointment export, field app log and out-of-hours desk record (extracts of "
               "5 October). Credits are those paid on these misses in the August, September and 2 October runs, "
               "landlord accounts included; the letting agents' accounts also carry credits for their properties "
               "outside Scotland, and only the credit for the property here is counted. Complaints are counted "
               "once each, including those the new complaints system carried over or reopened under a new "
               "reference in September.")
    p = doc.add_paragraph()
    r = p.add_run("Where to start: Parts & van stock misses by patch, July to September 2026")
    r.bold = True
    rows = [[pch, num(grid.get(pch, {}).get("Parts & van stock", 0))] for pch in PATCHES]
    rows.append(["Region", num(sum(grid.get(pch, {}).get("Parts & van stock", 0) for pch in PATCHES))])
    _table(doc, ["Patch", "Missed appointments"], rows, [6.0, 3.6], bold_last=True)
    _note(doc, "Patches as they stood on each appointment date; the boundaries changed on 1 August.")
    p = doc.add_paragraph()
    r = p.add_run("How the figures were built. ")
    r.bold = True
    p.add_run("A missed appointment is an attended appointment whose on-site time fell after its window (FS-07 "
              "1.1, 1.2). Each one was traced back along the engineer's day to its originating event (FS-07 3.1, "
              "3.2). Where that was a late log-on, the engineer's payroll number was matched to the out-of-hours "
              "desk record: every late log-on that followed a call-out clearing after 21:00 came between 11h00 "
              "and 11h04 after it. The same trace run over October 2025 to May 2026 finds no Rota & cover miss in "
              "any month. The weekly picture for the ops call uses reporting weeks Thursday to Wednesday; the "
              "worst week is the week of %s, %d missed appointments."
              % (D.fromisoformat(worst[0]).strftime("%-d %B"), worst[1]))
    cp = doc.core_properties
    cp.author = "Gillian Sanderson"
    cp.last_modified_by = "Gillian Sanderson"
    cp.created = dt.datetime(2026, 10, 6, 10, 12)
    cp.modified = dt.datetime(2026, 10, 6, 11, 40)
    cp.title = ""
    cp.comments = ""
    doc.save(path)


def main():
    OUT.mkdir(exist_ok=True)
    out, Gd, q3 = G.golden(ROOT / "target")
    fig = json.loads((Path(__file__).parent / "golden_figures.json").read_text())
    for k in ("family_counts", "credits_by_family", "complaints_by_family", "patch_family", "weekly"):
        assert json.loads(json.dumps(out[k], default=str)) == fig[k], k
    worst = chart(fig, OUT / "missed_appointments_q3_2026.png")
    assert worst[0] == fig["worst_week"]["start"] and worst[1] == fig["worst_week"]["count"]
    memo(fig, OUT / "missed_appointments_q3_2026.docx", worst)
    import seal
    seal.seal_ooxml(OUT / "missed_appointments_q3_2026.docx", "Gillian Sanderson", dt.datetime(2026, 10, 6, 11, 40))
    import writeup
    writeup.render()
    print("golden written", worst)


if __name__ == "__main__":
    main()
