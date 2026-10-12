"""The pack's documents. Short, functional, in the voice of the people who wrote them."""
import datetime as dt
from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib import colors

from world import MOVES, PATCH_NAME

COMPANY = "Kelvane Home Services"


def _pdf(path, title, blocks, meta_author, subject):
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=12.5,
                          alignment=TA_LEFT, spaceAfter=4)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=13, spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, spaceBefore=6,
                        spaceAfter=3)
    small = ParagraphStyle("s", parent=body, fontSize=8, textColor=colors.HexColor("#444444"))
    story = []
    for kind, val in blocks:
        if kind == "h1":
            story.append(Paragraph(val, h1))
        elif kind == "h2":
            story.append(Paragraph(val, h2))
        elif kind == "p":
            story.append(Paragraph(val, body))
        elif kind == "small":
            story.append(Paragraph(val, small))
        elif kind == "sp":
            story.append(Spacer(1, val * mm))
        elif kind == "table":
            t = Table(val, hAlign="LEFT")
            t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 8.5),
                                   ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.5),
                                   ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8E8E8"))]))
            story.append(t)
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                            bottomMargin=18 * mm, title=title, author=meta_author, subject=subject, creator="",
                            producer="")
    doc.build(story)


def procedure(path):
    fam = [["Family", "Charge here when"],
           ["Travel & routing", "the journey to an appointment took longer than the route plan allowed"],
           ["Job duration & standards", "a job took longer than its booked minutes"],
           ["Engineer attendance & punctuality", "the engineer logged on after the rostered start, or left a job "
                                                 "or the day early, for a reason within the engineer's own control"],
           ["Parts & van stock", "the engineer had to collect a part before going on to the next job"],
           ["Customer access", "the engineer was kept waiting on site for access"],
           ["Rota & cover", "the day could not be worked as planned because of how engineers were rostered or how "
                            "work was covered between engineers, teams and rotas"],
           ["Other", "none of the above (a vehicle breakdown, a road closed by the police)"]]
    fam = [[Paragraph(a, ParagraphStyle("t", fontName="Helvetica", fontSize=8.5, leading=10.5)),
            Paragraph(b, ParagraphStyle("t", fontName="Helvetica", fontSize=8.5, leading=10.5))] for a, b in fam]
    blocks = [
        ("small", COMPANY + " &nbsp;|&nbsp; Field Service Procedure FS-07 &nbsp;|&nbsp; Issue 3, effective "
                            "6 January 2025 &nbsp;|&nbsp; Owner: Head of Field Performance"),
        ("h1", "Missed appointment review"),
        ("p", "Applies to every field region. Replaces Issue 2 (March 2023)."),
        ("h2", "1. What counts"),
        ("p", "1.1 Customers book a repair or an annual service into a two-hour arrival window. An appointment is "
              "kept when the engineer's on-site time on the field app falls inside the window, and missed when "
              "the on-site time is after the window closes."),
        ("p", "1.2 The measure covers attended appointments: completed visits, and visits where the engineer arrived "
              "but could not get in. Appointments the customer cancelled or rebooked before the day are left out."),
        ("p", "1.3 Rates are missed appointments per attended appointment."),
        ("h2", "2. Cause families"),
        ("p", "Every missed appointment is charged to one family."),
        ("table", fam),
        ("h2", "3. Charging a missed appointment"),
        ("p", "3.1 A missed appointment is charged to its originating event: the earliest event without which the "
              "engineer would have reached the appointment inside its window."),
        ("p", "3.2 Lateness carried forward from earlier work is charged to the event that started it, not to the "
              "job the engineer was leaving when it showed."),
        ("p", "3.3 If two events would each have caused the miss on their own, the earlier is charged."),
        ("h2", "4. Verification"),
        ("p", "4.1 Each month Performance &amp; Quality draws 40 missed appointments at random from the previous "
              "month, checks the originating event against the field app and with the engineer, and records the "
              "family. The sample is tabled at the monthly field review."),
        ("p", "4.2 The sample is a check on how families are being charged. It is not a substitute for charging "
              "every missed appointment."),
        ("h2", "5. Reporting"),
        ("p", "5.1 Regions report missed appointments by family each month, for the region and for each patch."),
        ("p", "5.2 A family that accounts for the most missed appointments in a quarter is taken to the regional "
              "review with an owner."),
        ("sp", 4),
        ("small", "Issue 3 changes: 1.2 now counts no-access visits as attended; family 6 renamed from 'Resourcing'."),
    ]
    _pdf(path, "FS-07 Missed appointment review", blocks, "Field Performance", "Field Service Procedure FS-07")


def working_time_policy(path):
    blocks = [
        ("small", COMPANY + " &nbsp;|&nbsp; People Policy HR-FS-03 &nbsp;|&nbsp; Issue 5, April 2025"),
        ("h1", "Field working time"),
        ("p", "For field engineers and standby engineers on field contracts. Read with your contract of employment "
              "and the Working Time Regulations 1998."),
        ("h2", "1. Rostered hours"),
        ("p", "1.1 Day-roster engineers work Monday to Friday. The rostered start is 08:00, when you log on at home "
              "and set off for the first job. The field app records log-on, travel and on-site times."),
        ("p", "1.2 Standby engineers work the published standby rota, 17:00 to 08:00, and weekends by rota."),
        ("h2", "2. Breaks"),
        ("p", "2.1 A day over six hours carries a 30-minute unpaid break, taken between jobs."),
        ("h2", "3. Rest"),
        ("p", "3.1 Engineers must have at least eleven consecutive hours' rest between finishing one period of work "
              "and starting the next."),
        ("p", "3.2 Engineers must have at least one full day off in every seven, or two in every fourteen."),
        ("h2", "4. Weekly hours"),
        ("p", "4.1 Average weekly working time is limited to 48 hours over 17 weeks unless you have signed the "
              "opt-out. You can withdraw the opt-out with one month's notice."),
        ("h2", "5. Standby"),
        ("p", "5.1 Standby nights attract the standby allowance in the pay schedule, plus call-out pay for each "
              "attendance."),
        ("h2", "6. The call-out list"),
        ("p", "6.1 Day-roster engineers may join the voluntary call-out list, and leave it at any time."),
        ("p", "6.2 The out-of-hours desk may offer an attendance to an engineer on the list when no standby engineer "
              "can attend within an hour."),
        ("p", "6.3 Call-out pay is time and a half from leaving home to getting back, with a minimum of one hour. "
              "Time spent on a call-out counts as working time."),
        ("sp", 4),
        ("small", "Queries to People Services. This policy is reviewed every two years."),
    ]
    _pdf(path, "HR-FS-03 Field working time", blocks, "People Services", "People Policy HR-FS-03")


def standards(path):
    blocks = [
        ("small", COMPANY + " &nbsp;|&nbsp; Customer Service Standards, 2026 edition &nbsp;|&nbsp; extract: "
                            "sections 3 and 6"),
        ("h1", "Customer service standards (extract)"),
        ("h2", "3. Appointments"),
        ("p", "3.1 We give every booked visit a two-hour arrival window and call ahead if we are running late."),
        ("p", "3.2 If we arrive after the window closes we apply a goodwill credit of 30 pounds to the account that "
              "pays for the cover, or 50 pounds if the household is on our priority services register on the day "
              "of the visit."),
        ("p", "3.3 One credit is applied per household per visit day, however many of that day's appointments we "
              "missed."),
        ("p", "3.4 Credits are applied in the goodwill run on the second working day of the following month. We do "
              "not apply a credit to an account that has closed before the run."),
        ("p", "3.5 Other goodwill credits (a return visit we caused, a complaint we upheld) are applied in the same "
              "run under their own reason codes."),
        ("h2", "6. Complaints"),
        ("p", "6.1 We log a complaint against the account of the person who contacts us. A landlord can complain "
              "about a visit to a property they let; a tenant can complain about a visit to their home."),
        ("p", "6.2 When a closed complaint is reopened it keeps its reference and a new entry is logged."),
        ("p", "6.3 We acknowledge a complaint within two working days and aim to close it within eight weeks."),
        ("sp", 4),
        ("small", "Published to customers on the website. Internal reason codes: MA missed arrival window, RV return "
                  "visit, CG complaint goodwill."),
    ]
    _pdf(path, "Customer Service Standards 2026 (extract)", blocks, "Customer Experience",
         "Customer Service Standards")


def routing_note(path):
    blocks = [
        ("small", "Pathline Scheduler &nbsp;|&nbsp; Release 4.2 &nbsp;|&nbsp; note for field planning teams"),
        ("h1", "Release 4.2: what changes for field planning"),
        ("p", "Go-live: Wednesday 1 July 2026, all regions, for route plans built from that morning."),
        ("h2", "Road-speed model"),
        ("p", "4.2 replaces the road-speed tables with a model built from 2025 GPS traces. Motorway and trunk-road "
              "speeds are revised upward; urban and residential speeds are unchanged."),
        ("p", "In testing on 2025 diaries, planned travel on legs over about 16 km fell by around 15 per cent. Legs "
              "under that are planned as before."),
        ("h2", "Plan of record"),
        ("p", "The morning route plan remains the plan of record. Planned arrival and planned travel on the "
              "appointment are taken from it. Diaries are not re-planned during the day."),
        ("h2", "Other changes"),
        ("p", "Faster plan build (around 40 per cent); revised map tiles."),
        ("h2", "Known issues"),
        ("p", "PS-2291: planned travel can read one minute short where a leg crosses midnight. Fix in 4.2.1."),
        ("sp", 4),
        ("small", "Region champions: Scotland, Leslie Yates (dispatch). Questions to the planning systems team."),
    ]
    _pdf(path, "Pathline Scheduler 4.2 release note", blocks, "Planning Systems", "Release 4.2")


def _docx_props(doc, author, created):
    cp = doc.core_properties
    cp.author = author
    cp.last_modified_by = author
    cp.created = created
    cp.modified = created + dt.timedelta(minutes=37)
    cp.title = ""
    cp.comments = ""
    cp.keywords = ""
    cp.subject = ""


def review_notes(path):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    doc.add_heading("Scotland field review, 23 July 2026", level=1)
    doc.add_paragraph("Notes and actions. Teams call, 14:00 to 15:10.")
    doc.add_paragraph("Present: Gillian Sanderson (Field Operations Manager, chair), Leslie Yates (Dispatch Lead), "
                      "Bernard Akhtar (Technical Manager), Nigel Cook (Stores and Van Stock), Ryan Webster "
                      "(Resource Planning), Gemma Quinn (Performance and Quality). Lanarkshire and Fife team "
                      "leaders for item 3.")
    items = [
        ("1. Missed appointments",
         "Up sharply in the first three weeks of July against June. Leslie: Pathline 4.2 has cut the long legs and "
         "engineers are late into the afternoon. He would put the old speed tables back if the vendor will not "
         "change them. Action LY: raise with Pathline, ask for the 4.2 tables to be reviewed."),
        ("2. Service times",
         "Bernard: the April service checklist (flue gas analyser reading and CO check on every annual service) "
         "adds time and the 45-minute standard has not moved. Action BA: time a sample of services in August."),
        ("3. Log-ons",
         "Lanarkshire and Fife team leaders report more engineers logging on late this month. Gillian asked for a "
         "reminder on log-on times at next week's briefings. Action: team leaders."),
        ("4. Van stock",
         "Nigel: fortnightly van top-ups since the start of July mean more merchant runs mid-morning. Action NC: review the "
         "top-up list for diverter valves and PCBs."),
        ("5. Verification sample",
         "Gemma: the monthly sample stopped after May while she was on the 4.2 cut-over. Restarting in October. "
         "Action GQ."),
        ("6. Patch boundaries",
         "Ryan: the 1 August boundary changes go ahead as circulated. No change to team rosters."),
    ]
    for h, t in items:
        p = doc.add_paragraph()
        r = p.add_run(h)
        r.bold = True
        doc.add_paragraph(t)
    doc.add_paragraph("Next review: early October, with the quarter's figures.")
    _docx_props(doc, "Gillian Sanderson", dt.datetime(2026, 7, 23, 15, 22))
    doc.save(path)


def boundary_memo(path):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    doc.add_heading("Patch boundaries from Saturday 1 August 2026", level=1)
    doc.add_paragraph("From: Ryan Webster, Resource Planning. To: Scotland field managers and team leaders. "
                      "9 July 2026.")
    doc.add_paragraph("We are moving fourteen postcode districts between patches to even out the diaries ahead of the "
                      "autumn. Engineers stay in their teams; the districts below are planned by their new patch "
                      "team from the date of the change.")
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "District", "From", "To"
    for dist in ["G66", "G71", "G72", "G73", "G74", "G82", "EH30", "EH52", "EH54", "FK1", "ML5", "ML7", "PA3", "PA4"]:
        a, b = MOVES[dist]
        c = t.add_row().cells
        c[0].text, c[1].text, c[2].text = dist, PATCH_NAME[a], PATCH_NAME[b]
    doc.add_paragraph("")
    doc.add_paragraph("These changes take effect on Saturday 1 August 2026. Pathline picks up the new boundaries "
                      "overnight on 31 July. Our planning weeks run Thursday to Wednesday, so the change lands in the "
                      "middle of the planning week of Thursday 30 July to Wednesday 5 August: Thursday and Friday of "
                      "that week are planned on the old boundaries and the rest of it on the new ones.")
    doc.add_paragraph("Ryan")
    _docx_props(doc, "Ryan Webster", dt.datetime(2026, 7, 9, 11, 4))
    doc.save(path)


def conventions(path):
    txt = """# Field reporting conventions, Scotland region

Performance and Quality. Version 2.1, June 2026. Owner: Gemma Quinn.

These are the conventions behind the weekly ops call figures and the monthly field review.

- **Weeks.** Reporting weeks are the planning weeks, labelled by their first day.
- **Months and quarters.** Calendar months and quarters, by appointment date.
- **Patches.** An appointment counts in the patch its address belonged to on the appointment date.
- **Region.** Region figures are the six patches together.
- **Rates.** Per attended appointment, to one decimal place as a percentage.
- **Money.** Whole pounds in the weekly pack; pence in finance returns.

Questions to Gemma. Changes to these conventions go through the monthly field review.
"""
    Path(path).write_text(txt, encoding="utf-8")


def readme(path, counts):
    txt = """# Scotland field extracts: read me

Pulled on Monday 5 October 2026 by Performance and Quality for the Q3 field review. Internal to {co}.
Times are UK local time throughout.

## field_appointments_2025-07_2026-09.csv

Pathline Scheduler export. One row per booked appointment with an appointment date from 1 July 2025 to
30 September 2026 ({apps:,} rows).

| Field | Meaning |
|---|---|
| appointment_ref | appointment number |
| booked_on | date the appointment was booked |
| account_ref | customer account the appointment was booked on |
| postcode_district | postcode district of the property |
| patch | the property's patch as it stood when the appointment was booked |
| engineer_id | engineer the appointment was routed to |
| appointment_date, arrival_window | the booked two-hour arrival window |
| job_type, booked_minutes | job type and its standard minutes |
| planned_travel_min, planned_arrival | from the morning route plan of record |
| status | Completed, No access, Cancelled by customer, Rebooked by customer |
| en_route_at, on_site_at, job_started_at, off_site_at | field app times |
| travel_mode | handset travel status for the leg into the appointment: Direct, Via merchant, Vehicle recovery |
| window_met | Y when on_site_at falls inside the arrival window |

## engineer_day_log_2025-07_2026-09.csv

Field app. One row per day-roster engineer per working day ({daylog:,} rows): first log-on, last log-off and
the number of jobs routed.

## rota_2025-07_2026-09.csv

Roster system ({rota:,} rows). Day-roster engineers: D day roster, AL annual leave, TR training, SK sick.
Standby engineers: SB standby night, 17:00 to 08:00. rostered_start is given for D and SB.

## ooh_desk_callouts_2025-07_2026-10.xlsx

Out-of-hours desk platform. One row per emergency attendance, every night from 30 June 2025 to the night of
4 October 2026 ({ooh:,} rows). Cover is Standby rota, Call-out list (a day-roster engineer on the voluntary list)
or Contractor panel (the panel firm on the winter call-off). Attended by (payroll) is the payroll number of the
attending engineer; it is blank for contractor attendances.

## field_workforce_register.xlsx

HR register of the region's field and standby engineers on 5 October 2026, with leavers since 1 July 2025.

## customer_accounts_extract.csv

Customer accounts holding at least one property in the region ({accts:,} rows). account_type is Household or Landlord.
bill_payer_account: for a let property, the landlord account that pays for its cover. priority_register_from:
date the household joined the priority services register. merged_into and merged_on: duplicate accounts
merged in the August clean-up, with the account they were merged into.

## goodwill_credits_2025-08_2026-10.csv

Billing, company-wide (every region). Goodwill credits applied in the runs from August 2025 to 2 October 2026
({credits:,} rows). account_ref is the account credited. reason_code: MA missed arrival window, RV return
visit, CG complaint goodwill.

## goodwill_credit_runs.csv

Billing, company-wide. Count and value of each goodwill run as posted.

## customer_complaints_2025-07_2026-10.csv

Complaints system, company-wide (every region). One row per complaint entry received from 1 July 2025 to
5 October 2026
({complaints:,} rows). account_ref is the account the complaint is logged against. about_visit_date is the
visit the complaint concerns, where it concerns one. status is the entry's status: Open, Closed or Reopened.

## missed_appointment_verification_jan-may_2026.csv

Performance and Quality. The FS-07 section 4 verification sample, January to May 2026, 40 cases a month.

## postcode_district_patches.csv

Pathline's postcode district to patch lookup as it stands.

## Documents

- missed_appointment_review_procedure_FS-07.pdf: Field Service Procedure FS-07, Issue 3
- field_working_time_policy_HR-FS-03.pdf: People Policy HR-FS-03, Issue 5
- customer_service_standards_2026_extract.pdf: Customer Service Standards 2026, sections 3 and 6
- pathline_4.2_release_note.pdf: Pathline Scheduler release 4.2 note for field planning
- field_review_2026-07-23_notes.docx: notes and actions of the 23 July field review
- patch_boundaries_from_2026-08-01.docx: Resource Planning memo of 9 July on the 1 August boundaries
- field_reporting_conventions.md: Performance and Quality reporting conventions, version 2.1

## Sources

Every record and document in this folder is constructed for this exercise and describes a fictional
company; the people named are invented. Postcode districts are real outward codes for the places they
cover, used as labels only. Scottish bank holidays follow the published UK government list
(gov.uk/bank-holidays, Open Government Licence v3.0). Licence for the constructed material: CC BY 4.0.
""".format(co=COMPANY, **counts)
    Path(path).write_text(txt, encoding="utf-8")
