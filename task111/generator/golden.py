"""Golden figures, recomputed from the shipped files in target/ only."""
import csv
import datetime as dt
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

D = dt.date
MISSED = ("Missed appointment", "Appointments - arrival window")     # the category before and after the move
FAMILIES = ["Travel & routing", "Job duration & standards", "Engineer attendance & punctuality",
            "Parts & van stock", "Customer access", "Rota & cover", "Other"]
MOVED_BACK = {"G72": "Glasgow South & East", "G73": "Glasgow South & East", "G74": "Glasgow South & East",
              "G71": "Lanarkshire", "ML7": "Lanarkshire", "ML5": "Lanarkshire", "EH30": "Edinburgh & Lothians",
              "EH52": "Edinburgh & Lothians", "EH54": "Edinburgh & Lothians", "G66": "Fife & Forth Valley",
              "FK1": "Fife & Forth Valley", "PA4": "Ayrshire & Renfrewshire", "PA3": "Ayrshire & Renfrewshire",
              "G82": "Glasgow North & West"}
REBALANCE = D(2026, 8, 1)


def rows(T, name):
    with open(T / name, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def ts(s):
    return dt.datetime.strptime(s, "%Y-%m-%d %H:%M") if s else None


def mins_between(a, b):
    return int((b - a).total_seconds() // 60)


def load(T):
    G = {}
    G["apps"] = rows(T, "field_appointments_2025-07_2026-09.csv")
    G["daylog"] = {(r["engineer_id"], r["work_date"]): r for r in rows(T, "engineer_day_log_2025-07_2026-09.csv")}
    wb = load_workbook(T / "ooh_desk_callouts_2025-07_2026-10.xlsx", read_only=True)
    ws = wb["callouts"]
    it = ws.iter_rows(values_only=True)
    hdr = next(it)
    G["ooh"] = [dict(zip(hdr, r)) for r in it]
    wb = load_workbook(T / "field_workforce_register.xlsx", read_only=True)
    ws = wb["register"]
    it = ws.iter_rows(values_only=True)
    hdr = next(it)
    G["register"] = [dict(zip(hdr, r)) for r in it]
    G["accounts"] = {r["account_ref"]: r for r in rows(T, "customer_accounts_extract.csv")}
    G["credits"] = rows(T, "goodwill_credits_2025-08_2026-10.csv")
    G["complaints"] = rows(T, "customer_complaints_2025-07_2026-10.csv")
    G["map"] = {r["postcode_district"]: r["patch"] for r in rows(T, "postcode_district_patches.csv")}
    G["verified"] = rows(T, "missed_appointment_verification_jan-may_2026.csv")
    return G


def attended(a):
    return a["status"] in ("Completed", "No access")


def is_missed(a):
    if not attended(a):
        return False
    end = a["arrival_window"].split("-")[1]
    d = a["appointment_date"]
    return ts(a["on_site_at"]) > ts("%s %s" % (d, end))


def day_events(day_apps, logon, mode="trace"):
    """Delay events on one engineer-day, in time order: (time, kind, appointment_ref)."""
    ev = []
    d = day_apps[0]["appointment_date"]
    lo = mins_between(ts(d + " 08:00"), ts(logon))
    if lo > 5:
        ev.append((ts(logon), "logon", ""))
    for a in day_apps:
        if not attended(a):
            continue
        er, on, st, off = ts(a["en_route_at"]), ts(a["on_site_at"]), ts(a["job_started_at"]), ts(a["off_site_at"])
        trav = mins_between(er, on)
        if a["travel_mode"] == "Via merchant":
            ev.append((on, "parts", a["appointment_ref"]))
        elif a["travel_mode"] == "Vehicle recovery":
            ev.append((on, "other", a["appointment_ref"]))
        elif trav - int(a["planned_travel_min"]) > 14:
            ev.append((on, "travel", a["appointment_ref"]))
        if st is not None:
            if mins_between(on, st) > 10:
                ev.append((st, "access", a["appointment_ref"]))
            if mins_between(st, off) - int(a["booked_minutes"]) > 15:
                ev.append((off, "overrun", a["appointment_ref"]))
    ev.sort()
    return ev


KIND_FAMILY = {"travel": "Travel & routing", "overrun": "Job duration & standards",
               "logon": "Engineer attendance & punctuality", "parts": "Parts & van stock",
               "access": "Customer access", "other": "Other", "rest": "Rota & cover"}


def trace(G, d0, d1, cross_day=True):
    """Originating family of every missed appointment with appointment date in [d0, d1]."""
    by_day = defaultdict(list)
    for a in G["apps"]:
        if attended(a):
            by_day[(a["engineer_id"], a["appointment_date"])].append(a)
    pay = {r["Engineer ID"]: r["Payroll no"] for r in G["register"]}
    clears = defaultdict(list)
    for c in G["ooh"]:
        if c["Attended by (payroll)"]:
            clears[c["Attended by (payroll)"]].append(c["Cleared"])
    out = {}
    for (eid, d), apps in by_day.items():
        dd = D.fromisoformat(d)
        if not (d0 <= dd <= d1):
            continue
        apps.sort(key=lambda a: ts(a["en_route_at"]))
        missed = [a for a in apps if is_missed(a)]
        if not missed:
            continue
        logon = G["daylog"][(eid, d)]["logged_on_at"]
        ev = day_events(apps, logon)
        for m in missed:
            before = [e for e in ev if e[0] <= ts(m["on_site_at"])]
            assert before, "a miss with no upstream event %s" % m["appointment_ref"]
            kind = before[0][1]
            if kind == "logon" and cross_day:
                lo = ts(logon)
                for c in clears.get(pay[eid], []):
                    gap = mins_between(c, lo)
                    if 660 <= gap <= 664 and c.date() >= dd - dt.timedelta(days=1):
                        kind = "rest"
                        break
            out[m["appointment_ref"]] = (KIND_FAMILY[kind], m)
    return out


def one_hop(G, d0, d1):
    """Rung 1: each miss charged to the step immediately before it."""
    by_day = defaultdict(list)
    for a in G["apps"]:
        if attended(a):
            by_day[(a["engineer_id"], a["appointment_date"])].append(a)
    cnt = Counter()
    for (eid, d), apps in by_day.items():
        dd = D.fromisoformat(d)
        if not (d0 <= dd <= d1):
            continue
        apps.sort(key=lambda a: ts(a["en_route_at"]))
        logon = G["daylog"][(eid, d)]["logged_on_at"]
        for i, a in enumerate(apps):
            if not is_missed(a):
                continue
            er, on = ts(a["en_route_at"]), ts(a["on_site_at"])
            if a["travel_mode"] == "Via merchant":
                cnt["Parts & van stock"] += 1
            elif a["travel_mode"] == "Vehicle recovery":
                cnt["Other"] += 1
            elif mins_between(er, on) - int(a["planned_travel_min"]) > 14:
                cnt["Travel & routing"] += 1
            elif i == 0:
                cnt["Engineer attendance & punctuality"] += 1
            else:
                cnt["Job duration & standards"] += 1
    return cnt


def household_of(G, acct, visit, missed_hh_days):
    """Household (property account of record) behind an account on a visit date."""
    a = G["accounts"].get(acct)
    if a is None:
        return None
    if a["account_type"] == "Landlord":
        props = [k for k, v in G["accounts"].items() if v["bill_payer_account"] == acct]
        hits = [p for p in props if (p, visit) in missed_hh_days]
        return hits[0] if len(hits) == 1 else None
    if a["merged_into"]:
        return a["merged_into"]
    return acct


def golden(T):
    T = Path(T)
    G = load(T)
    q3 = trace(G, D(2026, 7, 1), D(2026, 9, 30))
    fam = Counter(f for f, _ in q3.values())
    out = {"q3_missed": len(q3), "family_counts": dict(fam)}
    # households-of-record per missed appointment
    def hh(acct):
        a = G["accounts"][acct]
        return a["merged_into"] or acct
    missed_hh_day = {}
    for ref, (f, a) in q3.items():
        missed_hh_day[(hh(a["account_ref"]), a["appointment_date"])] = f
    # landlord property index
    props_of = defaultdict(list)
    for k, v in G["accounts"].items():
        if v["bill_payer_account"]:
            props_of[v["bill_payer_account"]].append(k)

    def on_register(h, day):
        p = G["accounts"][h]["priority_register_from"]
        return bool(p) and p <= day

    def hits_of(acct, day):
        a = G["accounts"][acct]
        if a["account_type"] == "Landlord":
            return sorted(p for p in props_of[acct] if (p, day) in missed_hh_day)
        return [hh(acct)] if (hh(acct), day) in missed_hh_day else []
    # credits: one per household per visit day, to the account that pays for the cover (standards 3.2, 3.3). An
    # account carrying more lines for a visit day than it had households here missed that day is crediting its
    # properties in other regions as well: each household here takes the line at its own rate
    lines = defaultdict(list)
    for c in G["credits"]:
        if c["reason_code"] == "MA" and "2026-07-01" <= c["visit_date"] <= "2026-09-30" \
                and c["account_ref"] in G["accounts"]:
            lines[(c["account_ref"], c["visit_date"])].append(int(float(c["amount_gbp"])))
    cred = Counter()
    for (acct, day), amts in sorted(lines.items()):
        hits = hits_of(acct, day)
        if not hits:
            assert G["accounts"][acct]["account_type"] == "Landlord", (acct, day)
            continue
        if len(amts) == len(hits) == 1:
            cred[missed_hh_day[(hits[0], day)]] += amts[0]
            continue
        pool = list(amts)
        for h in hits:
            r = 50 if on_register(h, day) else 30
            assert r in pool, (acct, day, hits, amts)
            pool.remove(r)
            cred[missed_hh_day[(h, day)]] += r
    out["credits_by_family"] = dict(cred)
    # complaints: each complaint about a missed household-day counted once. A reopened complaint is the same
    # complaint, including one the new platform logged under a new reference; a complaint still open at the move
    # was taken across under a new reference, and the two records (same account, same moment received, same
    # visit) are one complaint
    first = {}
    for c in G["complaints"]:
        r = c["complaint_ref"]
        if r not in first or c["received_at"] < first[r]["received_at"]:
            first[r] = c
    canon, seen = {}, {}
    for r, c in sorted(first.items(), key=lambda x: (x[1]["received_at"], x[0])):
        canon[r] = seen.setdefault((c["account_ref"], c["received_at"], c["about_visit_date"]), r)
    comp = defaultdict(set)
    reopened_as_new = 0
    for c in G["complaints"]:
        if c["category"] not in MISSED or not ("2026-07-01" <= c["about_visit_date"] <= "2026-09-30"):
            continue
        acct = c["account_ref"]
        if acct not in G["accounts"]:
            continue
        hits = hits_of(acct, c["about_visit_date"])
        assert len(hits) <= 1, (acct, c["about_visit_date"], hits)
        if not hits:
            continue                                    # a complaint about a visit we kept
        key = (hits[0], c["about_visit_date"])
        if first[c["complaint_ref"]]["status"] == "Reopened":
            reopened_as_new += 1
            continue
        comp[missed_hh_day[key]].add(canon[c["complaint_ref"]])
    out["complaints_by_family"] = {k: len(v) for k, v in comp.items()}
    out["reopened_as_new"] = reopened_as_new
    # patch x family, structure in force on the appointment date
    grid = defaultdict(Counter)
    for ref, (f, a) in q3.items():
        dist = a["postcode_district"]
        dd = D.fromisoformat(a["appointment_date"])
        p = G["map"][dist] if (dd >= REBALANCE or dist not in MOVED_BACK) else MOVED_BACK[dist]
        grid[p][f] += 1
    out["patch_family"] = {p: dict(c) for p, c in grid.items()}
    # weekly series, reporting weeks Thursday to Wednesday, January to September 2026
    allm = trace(G, D(2025, 12, 25), D(2026, 9, 30))
    wk = defaultdict(Counter)
    for ref, (f, a) in allm.items():
        dd = D.fromisoformat(a["appointment_date"])
        start = dd - dt.timedelta(days=(dd.weekday() - 3) % 7)
        wk[start][f] += 1
    out["weekly"] = {k.isoformat(): dict(v) for k, v in sorted(wk.items())}
    out["one_hop_q3"] = dict(one_hop(G, D(2026, 7, 1), D(2026, 9, 30)))
    def in_post(day):
        n = 0
        for r in G["register"]:
            if r["Role"] != "Standby engineer":
                continue
            leave = r["Leave date"]
            if leave is None or leave.date() > day:
                n += 1
        return n
    out["standby_in_post"] = {"2025-07-01": in_post(D(2025, 7, 1)), "2026-06-01": in_post(D(2026, 6, 1)),
                              "2026-09-30": in_post(D(2026, 9, 30))}
    h1 = trace(G, D(2026, 1, 1), D(2026, 6, 30))
    out["h1_family_counts"] = dict(Counter(f for f, _ in h1.values()))
    att = lambda a0, a1: sum(1 for a in G["apps"] if attended(a) and a0 <= a["appointment_date"] <= a1)
    out["rates"] = {"h1_missed": len(h1), "h1_attended": att("2026-01-01", "2026-06-30"),
                    "q3_missed": len(q3), "q3_attended": att("2026-07-01", "2026-09-30")}
    same_day = trace(G, D(2026, 7, 1), D(2026, 9, 30), cross_day=False)
    out["same_day_q3"] = dict(Counter(f for f, _ in same_day.values()))
    out["forward"] = forward(G, fam, out["same_day_q3"])
    return out, G, q3


def night_of(t):
    """The out-of-hours night a timestamp belongs to (a night runs 17:00 to 08:00)."""
    return t.date() if t.hour >= 12 else t.date() - dt.timedelta(days=1)


def forward(G, fam, same_day):
    """Which Q3 causes carry into January to March, and the evidence from last winter."""
    fw = {}
    # the winter: Rota & cover by month under the cross-day trace, October 2025 to May 2026
    win = trace(G, D(2025, 10, 1), D(2026, 5, 31))
    rc = Counter(a["appointment_date"][:7] for f, a in win.values() if f == "Rota & cover")
    months = ["2025-10", "2025-11", "2025-12", "2026-01", "2026-02", "2026-03", "2026-04", "2026-05"]
    panel, latest, oct26, late_summer = Counter(), 0, 0, 0
    for c in G["ooh"]:
        n = night_of(c["Received"])
        m = n.isoformat()[:7]
        if c["Cover"] == "Contractor panel":
            if m in months:
                panel[m] += 1
            if D(2026, 10, 1) <= n <= D(2026, 10, 4):
                oct26 += 1
        if c["Cover"] == "Call-out list":
            clr = int((c["Cleared"] - dt.datetime(n.year, n.month, n.day)).total_seconds() // 60)
            if m in months:
                latest = max(latest, clr)
            if clr > 21 * 60 and n.month in (6, 7, 8, 9):
                late_summer += 1
    fw["winter"] = {"rota_cover_by_month": {m: rc.get(m, 0) for m in months},
                    "panel_by_month": {m: panel.get(m, 0) for m in months},
                    "latest_list_clear": "%02d:%02d" % (latest // 60, latest % 60),
                    "october_2026_panel": oct26, "summer_list_late": late_summer}
    # the causes that carry: every Q3 family except the one the winter panel removes
    carry = {f: n for f, n in fam.items() if f != "Rota & cover"}
    rank = sorted(carry.items(), key=lambda x: -x[1])
    fw["ranking"] = rank
    fw["answer"] = rank[0][0]
    # the forward-method grid
    q325 = Counter(f for f, _ in trace(G, D(2025, 7, 1), D(2025, 9, 30)).values())
    q126 = Counter(f for f, _ in trace(G, D(2026, 1, 1), D(2026, 3, 31)).values())
    seas = {f: fam.get(f, 0) * q126.get(f, 0) / q325[f] for f in FAMILIES if q325.get(f)}
    fw["grid"] = {"q3_carry_all": max(fam, key=fam.get),
                  "q3_carry_winter": rank[0][0],
                  "seasonal_ratio": max(seas, key=seas.get),
                  "last_winter_as_is": max(q126, key=q126.get),
                  "same_day_carry": max(same_day, key=same_day.get)}
    fw["grid_values"] = {"seasonal_ratio": {f: round(v, 1) for f, v in seas.items()},
                         "last_winter_q1_2026": dict(q126), "q3_2025": dict(q325)}
    # the load a solver carries forward: calls a night and standby in post
    nights = defaultdict(int)
    for c in G["ooh"]:
        nights[night_of(c["Received"])] += 1
    def per_night(a, b):
        ds = [d for d in nights if a <= d <= b]
        return round(sum(nights[d] for d in ds) / len(ds), 1)
    fw["calls_per_night"] = {"jan_2026": per_night(D(2026, 1, 1), D(2026, 1, 31)),
                             "q3_2026": per_night(D(2026, 7, 1), D(2026, 9, 30))}
    return fw
