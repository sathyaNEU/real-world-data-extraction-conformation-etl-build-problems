"""Write the data files of the shipped bundle."""
import csv
import datetime as dt
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font

from world import DISTRICTS_OLD, END, MOVES, PATCH_NAME, PATCHES, REBALANCE, START, at, patch_of

D = dt.date


def hm(m):
    m = int(m)
    return "%02d:%02d" % (m // 60, m % 60)


def stamp(d, m):
    return at(d, m).strftime("%Y-%m-%d %H:%M") if m is not None else ""


def wcsv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        for r in rows:
            w.writerow(r)


def write_data(W, T: Path, seed=111):
    rng = np.random.default_rng([seed, 8])
    T.mkdir(parents=True, exist_ok=True)
    files = {}

    # 1. appointments (spine), in booking-system order
    rows = []
    apps = sorted(W.appointments, key=lambda a: (a["date"], a["engineer_id"], a["window"][0],
                                                  a["arr_plan"] if a["arr_plan"] is not None else 9999,
                                                  a["appointment_ref"]))
    for a in apps:
        d = a["date"]
        attended = a["status"] in ("Completed", "No access")
        met = ""
        if attended:
            met = "Y" if a["arr"] <= a["window"][1] else "N"
        rows.append([
            a["appointment_ref"], a["booked_on"].isoformat(), a["account_ref"], a["district"],
            PATCH_NAME[a["stored_patch"]],
            a["engineer_id"], d.isoformat(), "%s-%s" % (hm(a["window"][0]), hm(a["window"][1])), a["job_type"],
            a["booked"], a["ptravel"] if attended else "", hm(a["arr_plan"]) if attended else "", a["status"],
            stamp(d, a["en_route"]) if attended else "", stamp(d, a["arr"]) if attended else "",
            stamp(d, a["start"]) if (attended and a["start"] is not None) else "",
            stamp(d, a["dep"]) if attended else "", a["travel_type"] if attended else "", met])
    f = "field_appointments_2025-07_2026-09.csv"
    wcsv(T / f, ["appointment_ref", "booked_on", "account_ref", "postcode_district", "patch", "engineer_id",
                 "appointment_date", "arrival_window", "job_type", "booked_minutes", "planned_travel_min",
                 "planned_arrival", "status", "en_route_at", "on_site_at", "job_started_at", "off_site_at",
                 "travel_mode", "window_met"], rows)
    files["appointments"] = f

    # 2. day-app shift log
    rows = []
    for r in sorted(W.day_records, key=lambda r: (r["date"], r["engineer_id"])):
        last = max(a["dep"] for a in r["acts"])
        off = last + int(rng.integers(18, 46))
        rows.append([r["engineer_id"], r["date"].isoformat(), stamp(r["date"], r["logon"]),
                     stamp(r["date"], off), len(r["jobs"])])
    f = "engineer_day_log_2025-07_2026-09.csv"
    wcsv(T / f, ["engineer_id", "work_date", "logged_on_at", "logged_off_at", "jobs_routed"], rows)
    files["daylog"] = f

    # 3. OOH desk call-out log (xlsx, the desk platform's export)
    wb = Workbook()
    ws = wb.active
    ws.title = "callouts"
    hdr = ["Call ref", "Received", "Emergency", "Postcode district", "Cover", "Attended by (payroll)",
           "Arrived", "Cleared", "Outcome"]
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True)
    seq = int(rng.integers(260000, 270000))
    outcomes = ["Made safe", "Repaired", "Made safe", "Repaired", "Referred to gas emergency service"]
    for c in sorted(W.calls, key=lambda c: (c["night"], c["received_m"])):
        seq += int(rng.integers(1, 5))
        pay = ""
        if c["engineer"]:
            pay = [e for e in W.engineers + W.standby if e["engineer_id"] == c["engineer"]][0]["payroll_no"]
        n = c["night"]
        ws.append(["OOH-%07d" % seq, at(n, c["received_m"]), c["category"], c["district"], c["cover"], pay,
                   at(n, c["arrived_m"]), at(n, c["cleared_m"]),
                   outcomes[int(rng.integers(len(outcomes)))] if c["category"] != "Leak" else "Made safe"])
    for col, wdt in zip("ABCDEFGHI", [13, 17, 15, 10, 16, 12, 17, 17, 30]):
        ws.column_dimensions[col].width = wdt
    for row in ws.iter_rows(min_row=2):
        for idx in (1, 6, 7):
            row[idx].number_format = "dd/mm/yyyy hh:mm"
    f = "ooh_desk_callouts_2025-07_2026-10.xlsx"
    wb.save(T / f)
    files["ooh"] = f

    # 4. workforce register
    wb = Workbook()
    ws = wb.active
    ws.title = "register"
    ws.append(["Engineer ID", "Payroll no", "Role", "Team patch", "Call-out list", "Start date", "Leave date"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for e in sorted(W.engineers + W.standby, key=lambda e: e["engineer_id"]):
        ws.append([e["engineer_id"], e["payroll_no"], e["role"], PATCH_NAME[e["home_patch"]],
                   ("Yes" if e.get("callout_list") else "No") if e["role"] == "Field engineer" else "n/a",
                   e.get("start_date") or D(int(rng.integers(2014, 2024)), int(rng.integers(1, 13)), 1),
                   e.get("end_date")])
    for row in ws.iter_rows(min_row=2):
        row[5].number_format = "dd/mm/yyyy"
        row[6].number_format = "dd/mm/yyyy"
    for col, wdt in zip("ABCDEFG", [11, 11, 17, 24, 12, 12, 12]):
        ws.column_dimensions[col].width = wdt
    f = "field_workforce_register.xlsx"
    wb.save(T / f)
    files["register"] = f

    # 5. rota: day codes and standby nights
    rows = []
    for e in W.engineers:
        for d in W.days:
            code = W.roster[e["engineer_id"]][d]
            rows.append([e["engineer_id"], d.isoformat(), code, "08:00" if code == "D" else ""])
    for s in W.standby:
        for n in W.nights:
            if n < START or n > END:
                continue
            if s["end_date"] and n >= s["end_date"]:
                continue
            if s["engineer_id"] in W.standby_on.get(n, []):
                rows.append([s["engineer_id"], n.isoformat(), "SB", "17:00"])
    rows.sort(key=lambda r: (r[1], r[0]))
    f = "rota_2025-07_2026-09.csv"
    wcsv(T / f, ["engineer_id", "rota_date", "duty", "rostered_start"], rows)
    files["rota"] = f

    # 6. customer accounts
    rows = []
    for h in W.households.values():
        rows.append([h["ref"], "Household", h["district"], h["opened"].isoformat(), h["plan"],
                     h["landlord"] or "", h["psr_from"].isoformat() if h["psr_from"] else "",
                     h["cancelled"].isoformat() if h["cancelled"] else "", "", ""])
        if h["dup"]:
            lo = max(h["opened"], D(2016, 1, 1)).toordinal()
            hi = D(2025, 6, 30).toordinal()
            dup_open = D.fromordinal(int(rng.integers(min(lo, hi - 30), hi)))
            rows.append([h["dup"], "Household", h["district"], dup_open.isoformat(), h["plan"],
                         h["landlord"] or "", "", "", h["ref"], "2026-08-18"])
    for l in W.landlords.values():
        rows.append([l["ref"], "Landlord", l["district"], l["opened"].isoformat(), "Landlord cover", "", "", "", "",
                     ""])
    rows.sort(key=lambda r: r[0])
    f = "customer_accounts_extract.csv"
    wcsv(T / f, ["account_ref", "account_type", "postcode_district", "opened_on", "plan", "bill_payer_account",
                 "priority_register_from", "cancelled_on", "merged_into", "merged_on"], rows)
    files["accounts"] = f

    # 7. goodwill credits ledger
    rows = [[c["credit_ref"], c["account_ref"], c["visit_date"].isoformat(), c["run_date"].isoformat(),
             c["reason"], "%d.00" % c["amount"]] for c in W.credits]
    f = "goodwill_credits_2025-08_2026-10.csv"
    wcsv(T / f, ["credit_ref", "account_ref", "visit_date", "credit_run_date", "reason_code", "amount_gbp"], rows)
    files["credits"] = f

    # 8. credit run summary (referee)
    runs = {}
    for c in W.credits:
        k = c["run_date"]
        runs.setdefault(k, [0, 0])
        runs[k][0] += 1
        runs[k][1] += c["amount"]
    rows = [[k.isoformat(), v[0], "%d.00" % v[1]] for k, v in sorted(runs.items())]
    f = "goodwill_credit_runs.csv"
    wcsv(T / f, ["credit_run_date", "credits", "total_gbp"], rows)
    files["runs"] = f

    # 9. complaints
    rows = [[c["complaint_ref"], c["account_ref"], c["received"].strftime("%Y-%m-%d %H:%M"), c["channel"],
             c["category"], c["about_visit_date"].isoformat() if c["about_visit_date"] else "", c["status"]]
            for c in W.complaints]
    f = "customer_complaints_2025-07_2026-10.csv"
    wcsv(T / f, ["complaint_ref", "account_ref", "received_at", "channel", "category", "about_visit_date",
                 "status"], rows)
    files["complaints"] = f

    # 10. postcode district map, current
    rows = []
    for dist in sorted(W.districts, key=lambda x: (x.rstrip("0123456789"), int(x.lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ")))):
        p = patch_of(dist, D(2026, 10, 5))
        rows.append([dist, PATCH_NAME[p]])
    f = "postcode_district_patches.csv"
    wcsv(T / f, ["postcode_district", "patch"], rows)
    files["map"] = f

    # 11. verified sample
    rows = []
    for i, v in enumerate(sorted(W.verified, key=lambda v: (v["a"]["date"], v["a"]["appointment_ref"]))):
        a = v["a"]
        rows.append(["MAV-26-%03d" % (i + 1), a["appointment_ref"], a["engineer_id"], a["date"].isoformat(),
                     "%s-%s" % (hm(a["window"][0]), hm(a["window"][1])), hm(v["origin_time"]), v["origin_ref"],
                     v["family"], v["note"]])
    f = "missed_appointment_verification_jan-may_2026.csv"
    wcsv(T / f, ["case_ref", "appointment_ref", "engineer_id", "appointment_date", "arrival_window",
                 "originating_event_time", "originating_appointment_ref", "cause_family", "verifier_note"], rows)
    files["verified"] = f
    return files
