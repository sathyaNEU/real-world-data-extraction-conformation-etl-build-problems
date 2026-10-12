"""Build task111's bundle: simulate, write target/, recompute the golden from target/, assert.

    python3 build.py            # from task111/generator
"""
import datetime as dt
import json
import sys
from collections import Counter
from pathlib import Path

import accounts
import days
import golden as G
import pack
import sample
import world

ROOT = Path(__file__).resolve().parent.parent
TARGET = ROOT / "target"
D = dt.date
SEED = 111


def simulate():
    W, rng = world.build_world(SEED)
    world.simulate_ooh(W, rng, {"seed": SEED})
    days.simulate_days(W, rng, {"seed": SEED})
    sample.make_twin(W, SEED)
    accounts.build_accounts(W, SEED)
    accounts.assign_households(W, SEED)
    accounts.portfolio_landlords(W, SEED)
    accounts.priority_joiners(W, SEED)
    accounts.credits_and_complaints(W, SEED)
    sample.verified_sample(W, SEED)
    return W


def write_docs(W):
    import csv
    import docs
    docs.procedure(TARGET / "missed_appointment_review_procedure_FS-07.pdf")
    docs.working_time_policy(TARGET / "field_working_time_policy_HR-FS-03.pdf")
    docs.standards(TARGET / "customer_service_standards_2026_extract.pdf")
    docs.routing_note(TARGET / "pathline_4.2_release_note.pdf")
    docs.review_notes(TARGET / "field_review_2026-07-23_notes.docx")
    docs.boundary_memo(TARGET / "patch_boundaries_from_2026-08-01.docx")
    docs.conventions(TARGET / "field_reporting_conventions.md")

    def n(f):
        with open(TARGET / f, encoding="utf-8") as fh:
            return sum(1 for _ in fh) - 1
    from openpyxl import load_workbook
    ooh = load_workbook(TARGET / "ooh_desk_callouts_2025-07_2026-10.xlsx", read_only=True)["callouts"].max_row - 1
    docs.readme(TARGET / "scotland_field_extracts_readme.md",
                {"apps": n("field_appointments_2025-07_2026-09.csv"), "daylog": n("engineer_day_log_2025-07_2026-09.csv"),
                 "rota": n("rota_2025-07_2026-09.csv"), "ooh": ooh, "accts": n("customer_accounts_extract.csv"),
                 "credits": n("goodwill_credits_2025-08_2026-10.csv"),
                 "complaints": n("customer_complaints_2025-07_2026-10.csv")})


def main():
    W = simulate()
    # the bundle is regenerated whole: clear the previous cut so a renamed file cannot linger
    if TARGET.exists():
        for p in TARGET.iterdir():
            if p.is_file():
                p.unlink()
    files = pack.write_data(W, TARGET, SEED)
    write_docs(W)
    import seal
    seal.seal(TARGET)
    out, Gd, q3 = G.golden(TARGET)
    import asserts
    n = asserts.run(W, out, Gd, q3, TARGET)
    fig = {k: v for k, v in out.items()}
    fig["assertions"] = n
    (Path(__file__).parent / "golden_figures.json").write_text(json.dumps(fig, indent=1, sort_keys=True,
                                                                          default=str))
    print("assertions green:", n)
    print(json.dumps({k: out[k] for k in ("family_counts", "credits_by_family", "complaints_by_family",
                                          "same_day_q3", "one_hop_q3", "worst_week", "corpus")}, indent=1))
    return W, out


if __name__ == "__main__":
    main()
