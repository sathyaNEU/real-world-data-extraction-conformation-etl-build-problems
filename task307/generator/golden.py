"""The goldens: mampo_tip2028_project_ranking.xlsx (the scoring and the count-program tables) and
mampo_tip2028_executive_brief.pptx (the recommendation and the chart of the four annual benefits), drawn from
figures.json.

    python3 task307/generator/golden.py <golden dir>
"""
import datetime as dt
import io
import json
import os
import re
import sys
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import params as P  # noqa: E402
import postprocess  # noqa: E402
import records as R  # noqa: E402

DECK = "mampo_tip2028_executive_brief.pptx"
XLSX = "mampo_tip2028_project_ranking.xlsx"
AUTHOR = P.PEOPLE["lead"]
ORDER = ["TIP-28-014", "TIP-28-021", "TIP-28-009", "TIP-28-030"]
SHORT = {"TIP-28-014": "Route 9 auxiliary lane", "TIP-28-021": "Ferry Street exit ramp",
         "TIP-28-009": "Center Street Bridge", "TIP-28-030": "Southport signals and turn lanes"}
LOC = {s[0]: s[1] for s in R.STATIONS}


def c(n):
    return "{:,}".format(int(n))


def _text(slide, left, top, width, height, text, size, bold=False, colour=None):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tb.text_frame.word_wrap = True
    tb.text_frame.text = text
    for para in tb.text_frame.paragraphs:
        for run in para.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            if colour:
                run.font.color.rgb = RGBColor.from_string(colour)
    return tb


def _title(slide, text):
    t = slide.shapes.title
    t.text = text
    t.left, t.top, t.width, t.height = Inches(0.45), Inches(0.3), Inches(12.4), Inches(0.95)
    p = t.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.runs[0].font.size = Pt(25)
    p.runs[0].font.bold = True


def deck(F, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[5])
    aw = F["award"]
    _title(s, "Put the Center Street Bridge widening forward: about %s vehicle-hours a year saved"
           % c(F["annual_vh"][aw]))
    order = sorted(ORDER, key=lambda k: F["annual_vh"][k])
    cd = CategoryChartData()
    cd.categories = [SHORT[k] for k in order]
    cd.add_series("Vehicle-hours saved a year (thousands)", [F["annual_vh"][k] / 1000.0 for k in order])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.45), Inches(1.35), Inches(7.9), Inches(5.2), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = ("Annual benefit, 2028 AM and PM peak hours re-assigned on each build network "
                                      "(thousand vehicle-hours a year)")
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    va = ch.value_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = 0, 350, 50
    va.has_major_gridlines = False
    va.tick_labels.font.size = Pt(10)
    ch.category_axis.tick_labels.font.size = Pt(11)
    pl = ch.plots[0]
    pl.gap_width = 70
    ser = pl.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = RGBColor(0xA9, 0xB1, 0xB9)
    for i, k in enumerate(order):
        pt = ser.points[i]
        if k == aw:
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = RGBColor(0x1D, 0x5C, 0x7A)
        dl = pt.data_label
        dl.has_text_frame = True
        dl.text_frame.text = c(F["annual_vh"][k])
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.text_frame.paragraphs[0].runs[0].font.size = Pt(10.5)
        dl.text_frame.paragraphs[0].runs[0].font.bold = k == aw
    ru = F["runner_up"]
    pm = F["per_musd_vh"]
    _text(s, 8.65, 1.45, 4.25, 5.1,
          "Center Street saves %s vehicle-hours a year, %s for each million dollars of its USD %.1f million cost.\n\n"
          "The Route 9 auxiliary lane is next at %s (%s per million). Re-timed at today's loaded volumes it looks "
          "like the biggest saver, but once the 2028 trips re-route, the extra lane at Millbrook only "
          "delivers more traffic to the Route 9 bridge, which then sets the queue, and the new Center Street lane "
          "draws trips off Route 9 as well.\n\nSouthport signals and turn lanes: %s (%s per million). Ferry Street "
          "exit ramp: %s (%s per million)."
          % (c(F["annual_vh"][aw]), c(pm[aw]), F["cost_musd"][aw], c(F["annual_vh"][ru]), c(pm[ru]),
             c(F["annual_vh"]["TIP-28-030"]), c(pm["TIP-28-030"]),
             c(F["annual_vh"]["TIP-28-021"]), c(pm["TIP-28-021"])), 13)
    _text(s, 0.45, 6.8, 12.4, 0.4,
          "TDM v4.2 subarea network and 2028 trip tables, user equilibrium to a relative gap below 0.000001. Annual "
          "benefit per the TIP scoring guidance: peak-hour savings times 247 weekdays times 2.9. Travel Forecasting "
          "Program, 23 October 2026.", 9, colour="6B6B6B")


    prs.core_properties.author = AUTHOR
    prs.core_properties.last_modified_by = AUTHOR
    prs.core_properties.title = "FY2028 TIP congestion-relief slot: executive brief"
    prs.core_properties.comments = ""
    prs.core_properties.created = dt.datetime(2026, 10, 21, 15, 10)
    prs.core_properties.modified = dt.datetime(2026, 10, 23, 11, 42)
    prs.core_properties.revision = 6
    prs.save(path)
    _fix_embedded(path)
    postprocess.scrub(path, P.PRODUCER, "2026-10-23")


def _fix_embedded(path):
    """the chart workbook python-pptx embeds carries the build clock: give it the deck's own date and fixed entry
    stamps so two builds are byte-identical"""
    with zipfile.ZipFile(path) as z:
        items = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    out = []
    for n, b in items:
        if n.endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(b)) as zz:
                inner = [(i.filename, zz.read(i.filename)) for i in zz.infolist()]
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zo:
                for m, bb in inner:
                    if m == "docProps/core.xml":
                        bb = re.sub(rb"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", b"2026-10-23T11:42:00Z", bb)
                        bb = re.sub(rb"<dc:creator>[^<]*</dc:creator>", b"<dc:creator>" + AUTHOR.encode() +
                                    b"</dc:creator>", bb)
                    zi = zipfile.ZipInfo(m, date_time=(1980, 1, 1, 0, 0, 0))
                    zi.compress_type = zipfile.ZIP_DEFLATED
                    zo.writestr(zi, bb)
            b = buf.getvalue()
        if n.startswith("ppt/slideMasters/") and n.endswith(".xml"):
            b = b.replace("\u2013".encode(), "\u2022".encode())
        out.append((n, b))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in out:
            zi = zipfile.ZipInfo(n, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, b)


def _head(ws, row):
    hdr, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1D5C7A")
    for cell in ws[row]:
        cell.font, cell.fill = hdr, fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def workbook(F, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "TIP scoring"
    ws.append(["FY2028 TIP congestion-relief candidates: annual benefit under the TIP scoring guidance (May 2026)"])
    ws["A1"].font = Font(bold=True)
    ws.append(["TIP ID", "Project", "Cost, USD M (YOE)", "Annual benefit, vehicle-hours", "Benefit per USD 1M, vehicle-hours",
               "Rank"])
    _head(ws, 2)
    for rank, k in enumerate(F["ranking"], start=1):
        ws.append([k, F["names"][k], F["cost_musd"][k], F["annual_vh"][k], F["per_musd_vh"][k], rank])
        ws.cell(row=ws.max_row, column=3).number_format = "0.0"
        ws.cell(row=ws.max_row, column=4).number_format = "#,##0"
        ws.cell(row=ws.max_row, column=5).number_format = "#,##0"
    ws.append([])
    ws.append(["Recommended for the open slot: %s (%s)." % (F["award_name"], F["award"])])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append(["Benefit: reduction in 2028 AM plus PM peak-hour vehicle-hours with each project coded into the v4.2 "
               "subarea network and the trip tables re-assigned to user equilibrium, times 247 weekdays times 2.9."])
    for col, w in zip("ABCDEF", (12, 46, 16, 20, 20, 8)):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[2].height = 32

    ws2 = wb.create_sheet("Count stations 2025")
    ws2.append(["Permanent count stations: 2025 AADT and K-factor (Traffic Count Program Manual, section 4)"])
    ws2["A1"].font = Font(bold=True)
    ws2.append(["Station", "Location", "AADT, vehicles", "K-factor"])
    _head(ws2, 2)
    for st in sorted(F["aadt"]):
        ws2.append([st, LOC[st], F["aadt"][st], float(F["k_factor"][st])])
        ws2.cell(row=ws2.max_row, column=3).number_format = "#,##0"
        ws2.cell(row=ws2.max_row, column=4).number_format = "0.000"
    ws2.append([])
    ws2.append(["S-0103 includes its replacement S-0103R from 15 July 2025. Days on which a loop or detector card was "
                "faulty, up to its repair, are not valid days; nor is an hour that carries a missed poll's vehicles."])
    for col, w in zip("ABCD", (10, 42, 16, 10)):
        ws2.column_dimensions[col].width = w

    ws3 = wb.create_sheet("Base-year validation")
    ws3.append(["TDM v4.2 2025 base year against the river-bridge count stations (model documentation 6.4)"])
    ws3["A1"].font = Font(bold=True)
    ws3.append(["Station", "Location", "Direction", "Peak hour", "Model less count, % of count"])
    _head(ws3, 2)
    for key, v in F["validation_pct"].items():
        st, d, per = key.split()
        ws3.append([st, LOC[st], d, per, float(v)])
        ws3.cell(row=ws3.max_row, column=5).number_format = "0.0"
    ws3.append([])
    ws3.append(["Model volumes converted from passenger-car equivalents to vehicles; counts averaged over valid "
                "Tuesdays to Thursdays, March to May and September to October 2025."])
    for col, w in zip("ABCDE", (10, 30, 10, 10, 20)):
        ws3.column_dimensions[col].width = w
    ws3.row_dimensions[2].height = 32

    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.properties.created = dt.datetime(2026, 10, 20, 9, 5)
    wb.properties.modified = dt.datetime(2026, 10, 23, 10, 58)
    wb.save(path)
    postprocess.scrub(path, P.PRODUCER, "2026-10-23")
    postprocess.finish_office(path, "2026-10-20T14:05:00Z", "2026-10-23T15:58:00Z", P.ORG, AUTHOR)


def write(golden_dir):
    F = json.load(open(os.path.join(HERE, "figures.json")))
    os.makedirs(golden_dir, exist_ok=True)
    deck(F, os.path.join(golden_dir, DECK))
    workbook(F, os.path.join(golden_dir, XLSX))


if __name__ == "__main__":
    write(sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(HERE), "golden"))
