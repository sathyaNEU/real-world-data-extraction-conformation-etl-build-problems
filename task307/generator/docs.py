"""task307 documents: the model documentation excerpt, the TIP scoring guidance, the count program manual and the CMP
reliability measures (pdf); the TIP candidate workbook, the count station inventory and the crash summary (xlsx)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import net
import params as P
import postprocess


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Times-Roman", fontSize=10.2, leading=13.6, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=13.5, leading=17, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.4, leading=14,
                        spaceBefore=8, spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontName="Helvetica", fontSize=8.2, leading=10.8,
                           textColor=colors.HexColor("#444444"))
    return body, h1, h2, small


def _footer(text):
    def foot(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.4)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(19 * mm, 10 * mm, text)
        canvas.drawRightString(letter[0] - 19 * mm, 10 * mm, "Page %d" % d.page)
        canvas.restoreState()
    return foot


def _table(rows, widths, size=8.6):
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=size, leading=size + 2.2)
    head = ParagraphStyle("head", parent=cell, fontName="Helvetica-Bold")
    rows = [[Paragraph(str(v), head if r == 0 else cell) for v in row] for r, row in enumerate(rows)]
    t = Table(rows, colWidths=[w * mm for w in widths], repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9A9A9A")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6EAF0")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


SAVED = {"2026-03-18": ("091233", "143810", "-05'00'"), "2026-05-21": ("101102", "161944", "-05'00'"),
         "2024-02-07": ("083512", "112906", "-06'00'"), "2026-02-12": ("135041", "152218", "-06'00'")}


def _pdf(path, title, subject, story, foot, stamp):
    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=19 * mm, rightMargin=19 * mm, topMargin=17 * mm,
                            bottomMargin=18 * mm, title=title, author=P.PRODUCER, creator=P.PRODUCER, subject=subject,
                            invariant=1)
    doc.build(story, onFirstPage=_footer(foot), onLaterPages=_footer(foot))
    postprocess.scrub(path, P.PRODUCER, stamp)
    postprocess.pdf_clock(path, *SAVED[stamp])


def model_doc(path):
    body, h1, h2, small = _styles()
    vdf = [["FTYPE", "Facility", "Capacity per lane (vehicles per hour)", "alpha", "beta"]]
    names = {1: "Freeway", 12: "Freeway, lane-drop or weaving section", 2: "Major river bridge", 3: "Ramp",
             4: "Arterial", 5: "Collector or downtown grid", 9: "Centroid connector"}
    key = {1: "freeway", 12: "freeway", 2: "bridge", 3: "ramp", 4: "arterial", 5: "collector", 9: "connector"}
    vkey = {1: "freeway", 12: "bridge", 2: "bridge", 3: "ramp", 4: "arterial", 5: "collector", 9: "connector"}
    for f in (1, 12, 2, 3, 4, 5, 9):
        a, b = net.VDF[vkey[f]]
        cap = "not constrained" if f == 9 else "{:,}".format(net.LANE_CAP[key[f]])
        vdf.append([f, names[f], cap, "%.2f" % a if f != 9 else "0", "%.0f" % b if f != 9 else "-"])
    s = [Paragraph("Merrimont Regional Travel Demand Model, version 4.2", h1),
         Paragraph("Model documentation, chapters 4 to 6 (network, volume-delay functions, assignment). Travel "
                   "Forecasting Program, %s. March 2026." % P.ORG, small),
         Paragraph("4. Highway network", h2),
         Paragraph("4.1 The highway network is held as two-way link records. Each record carries the link's A and B "
                   "nodes, the number of lanes in each direction (AB_LANES from A to B, BA_LANES from B to A; a lane "
                   "count of zero closes that direction), the length in miles, the free-flow speed in miles per hour "
                   "and the facility type (FTYPE). A link's free-flow time is its length divided by its free-flow "
                   "speed, and its capacity in a direction is that direction's lanes times the capacity per lane of "
                   "its facility type in Table 4.2.", body),
         Paragraph("4.2 Zones are numbered by district: 100s on the west bank, 200s downtown, 300s east, 400s the "
                   "port. Numbers retired at the 2023 zone split are not reused. Each zone loads onto the network "
                   "through its centroid connector (FTYPE 9), which is not capacity constrained.", body),
         Paragraph("4.3 Subarea extracts. Project appraisal in the Ashkettle crossings corridor uses the subarea "
                   "network extracted from the regional network with its trip tables, so that a run takes seconds "
                   "rather than hours. The extract keeps every link a corridor trip can use; trips with an end "
                   "outside the subarea enter and leave through the extract's boundary zones, which carry ordinary "
                   "zone numbers.", body),
         Paragraph("5. Volume-delay functions", h2),
         Paragraph("5.1 A link's travel time at volume v is its free-flow time t0 times (1 + alpha (v/c)^beta), "
                   "where c is its capacity in that direction and alpha and beta are those of its facility type "
                   "(Table 4.2). Lane-drop and weaving sections of freeway (FTYPE 12) take the bridge curve, which "
                   "the 2023 recalibration fitted to the corridor's queue discharge data.", body),
         _table(vdf, [16, 58, 50, 18, 16]),
         Paragraph("Table 4.2. Capacities and volume-delay parameters by facility type.", small),
         Paragraph("6. Assignment", h2),
         Paragraph("6.1 The AM and PM peak-hour trip tables are assigned separately. Assignment is to user "
                   "equilibrium: every trip in the table uses a route whose travel time is no greater than that of "
                   "any other route available to it between its zones, at the link times that all the trips together "
                   "produce. Production runs converge to a relative gap below 0.000001.", body),
         Paragraph("6.2 Peak hours. The AM peak hour is 7:00 to 8:00 a.m. and the PM peak hour 5:00 to 6:00 p.m., "
                   "local time.", body),
         Paragraph("6.3 Heavy vehicles. The trip tables are in vehicles. Assigned volumes in the station reports are "
                   "stated in passenger-car equivalents, a heavy vehicle counting as 2.0, with the model's "
                   "heavy-vehicle percentage for the link alongside.", body),
         Paragraph("6.4 Validation. The 2025 base year is checked against the count program's permanent stations on "
                   "the river crossings and on Route 9. "
                   "The comparison is in vehicles: the model's peak-hour volume against the average count for the "
                   "same hour on the valid Tuesdays, Wednesdays and Thursdays of March to May and September to "
                   "October 2025, federal holidays excluded, a valid day being as the Traffic Count Program Manual "
                   "defines it. The check is reported as the model's volume less the count, as a percentage of the "
                   "count, by station, direction and peak hour.", body)]
    _pdf(path, "Merrimont Regional Travel Demand Model v4.2: documentation, chapters 4 to 6",
         "Travel demand model documentation", s, "%s  |  TDM v4.2 documentation" % P.ORG, "2026-03-18")


def guidance(path):
    body, h1, h2, small = _styles()
    s = [Paragraph("Transportation Improvement Program: project scoring guidance", h1),
         Paragraph("Adopted by the Technical Advisory Committee, May 2026. Applies to projects proposed for the "
                   "FY2028 to FY2031 TIP.", small),
         Paragraph("1. Scope", h2),
         Paragraph("This guidance sets how candidates for the TIP's congestion-relief category are scored. Safety, "
                   "transit and maintenance candidates are scored under their own sections and are not covered "
                   "here.", body),
         Paragraph("2. Congestion-relief benefit", h2),
         Paragraph("2.1 A candidate's benefit is the reduction, with the project open in its opening year, in the "
                   "total vehicle-hours that the region's AM and PM peak-hour trips spend travelling on the "
                   "network, against the opening-year no-build network. Both are taken from the current regional "
                   "model, on the network the project is coded into and the opening-year trip tables. Trips are "
                   "taken as the trip tables give them; the category does not credit trips a project would add or "
                   "remove.", body),
         Paragraph("2.2 The annual benefit is the AM peak-hour reduction plus the PM peak-hour reduction, times %d "
                   "typical weekdays (the non-holiday weekdays outside the week of 25 December) times %.1f, the "
                   "number of hours of the peak periods that one peak hour stands for." % (P.WEEKDAYS, P.PEAK_EXPANSION), body),
         Paragraph("2.3 Each candidate is reported with its annual benefit and its annual benefit per million "
                   "dollars of year-of-expenditure cost.", body),
         Paragraph("3. Award", h2),
         Paragraph("3.1 When one congestion-relief slot is open and every candidate fits its funding, the slot goes "
                   "to the candidate with the largest annual benefit. The executive team's recommendation to the "
                   "Policy Board names the candidate, its annual benefit and its benefit per million dollars, and "
                   "the next candidate's.", body),
         Paragraph("4. Records", h2),
         Paragraph("Candidates are coded by the Travel Forecasting Program from the sponsor's project sheet. The "
                   "coding sheet in the candidate workbook is the record of what each candidate changes on the "
                   "network.", body)]
    _pdf(path, "TIP project scoring guidance (adopted May 2026)", "TIP scoring guidance", s,
         "%s  |  TIP scoring guidance  |  May 2026" % P.ORG, "2026-05-21")


def count_manual(path):
    body, h1, h2, small = _styles()
    s = [Paragraph("Traffic Count Program Manual", h1),
         Paragraph("%s, Data Services. Edition of February 2024." % P.ORG, small),
         Paragraph("3. Permanent count stations", h2),
         Paragraph("3.1 Each permanent station reports an hourly volume in each direction. Hours are stamped by "
                   "hour ending, in local time, as the station clock keeps it; on the day the clocks go forward the "
                   "day has 23 hours, and on the day they go back it has 25, the repeated hour appearing twice.", body),
         Paragraph("3.2 Each hourly record carries a status. OK means the station answered the server's poll for the "
                   "hour, and the volume is what its counters returned. NR (no response) means the station did not "
                   "answer, and the server writes a volume of zero. EST means the vendor estimated the volume from "
                   "adjacent days; an estimate is not a count. When a poll is missed while the station is running, "
                   "across midnight as at any other hour, the station's counters keep counting, and the next answered "
                   "poll returns everything counted since "
                   "the last answered one; after an outage of whole days the counters restart, and the first poll "
                   "back is that hour's count.", body),
         Paragraph("3.3 When a station's sensors are replaced at the same location, the replacement takes a new "
                   "station number with the suffix R and the inventory records the change. The two numbers are one "
                   "station's series for the year.", body),
         Paragraph("4. Annual statistics", h2),
         Paragraph("4.1 A valid day for a station is a day on which every hour the day has is, in both directions, a "
                   "true count of that hour's traffic.", body),
         Paragraph("4.2 AADT is computed by the AASHTO method from valid days only: for each month, the two-way "
                   "daily volume is averaged over the valid days of each day of the week; the month's average daily "
                   "traffic is the average of its seven day-of-week averages; AADT is the average of the twelve "
                   "monthly averages. AADT is reported in whole vehicles.", body),
         Paragraph("4.3 The K-factor is the 30th-highest two-way hourly volume of the year, over the hours that are a "
                   "true count of the hour's traffic in both directions, divided by the year's AADT, reported to three "
                   "decimals.", body)]
    _pdf(path, "Traffic Count Program Manual (February 2024)", "Count program manual", s,
         "%s  |  Traffic Count Program Manual  |  2024" % P.ORG, "2024-02-07")


def cmp_doc(path):
    body, h1, h2, small = _styles()
    s = [Paragraph("Congestion Management Process: travel time reliability measures", h1),
         Paragraph("%s. Technical memorandum, February 2026." % P.ORG, small),
         Paragraph("1. Corridors and data", h2),
         Paragraph("The CMP monitors four corridors from probe travel times purchased by segment and 15-minute "
                   "interval. The vendor's measurement_tstamp is in UTC and marks the end of the 15-minute interval. "
                   "The segment reference file lists each segment's corridor, direction, order, length "
                   "and reference speed. A corridor's AM measure is taken in its inbound direction (northbound or "
                   "eastbound) and its PM measure in its outbound direction.", body),
         Paragraph("2. Planning time index", h2),
         Paragraph("2.1 A corridor's travel time for a 15-minute interval is the sum of its segments' travel times "
                   "for that interval. An interval is used only when every segment of the corridor reports.", body),
         Paragraph("2.2 The planning time index for a peak hour is the 95th-percentile corridor travel time over "
                   "the year's weekday intervals that start within that hour (the AM peak hour 7:00 to 8:00 a.m., "
                   "the PM peak hour 5:00 to 6:00 p.m., local time, which in Merrimont is US Central time), federal holidays excluded, divided by the "
                   "corridor's free-flow travel time. The 95th percentile is the nearest-rank value: the interval "
                   "times sorted from shortest, the value at rank ceil(0.95 n). The free-flow travel time is the sum "
                   "of the segments' lengths divided by their reference speeds. The index is reported to two "
                   "decimals.", body)]
    _pdf(path, "CMP travel time reliability measures (February 2026)", "CMP technical memorandum", s,
         "%s  |  Congestion Management Process  |  Feb 2026" % P.ORG, "2026-02-12")


def _sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DCE3EA")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w


def tip_workbook(path, coding):
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidates"
    rows = [[p["tip_id"], p["name"], s, p["cost_m"], d] for p, s, d in coding["summary"]]
    _sheet(ws, ["TIP ID", "Project", "Sponsor", "Cost, $M (YOE)", "Description"], rows, [12, 40, 22, 14, 70])
    ws2 = wb.create_sheet("Model coding")
    _sheet(ws2, ["TIP ID", "LINK_ID", "A_NODE", "B_NODE", "Change", "From", "To"], coding["rows"], [12, 9, 9, 9, 34, 10, 10])
    wb.save(path)
    postprocess.scrub(path, P.PRODUCER, "2026-09-30")
    postprocess.finish_office(path, "2026-08-18T14:02:11Z", "2026-09-30T16:40:27Z", P.ORG, P.PEOPLE["lead"])


def station_inventory(path, rows, maint):
    wb = Workbook()
    ws = wb.active
    ws.title = "Stations"
    _sheet(ws, ["Station", "Location", "Directions", "Type", "In service", "Status", "Notes"], rows,
           [11, 40, 11, 16, 12, 12, 60])
    ws2 = wb.create_sheet("Maintenance 2025")
    _sheet(ws2, ["Station", "Completed", "Work order", "Work done"], maint, [11, 12, 12, 90])
    wb.save(path)
    postprocess.scrub(path, P.PRODUCER, "2026-01-09")
    postprocess.finish_office(path, "2025-01-06T15:21:40Z", "2026-01-09T10:12:03Z", P.ORG, "Data Services")


def crash_workbook(path, df):
    wb = Workbook()
    ws = wb.active
    ws.title = "By segment"
    _sheet(ws, list(df.columns), df.values.tolist(), [40, 8, 9, 9, 8, 8])
    wb.save(path)
    postprocess.scrub(path, P.PRODUCER, "2026-04-02")
    postprocess.finish_office(path, "2026-03-30T13:10:00Z", "2026-04-02T09:44:51Z", P.ORG, "Safety Program")
