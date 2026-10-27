"""The main call computed from the cut pack alone, the way the model documentation and the scoring guidance file it:
the subarea extract's two-way link records, its 2028 trip tables and each candidate's coding, assigned to user
equilibrium for the no-build network and for each build network, and the reduction in peak-hour vehicle-hours
annualised. Also the cheaper routes a response takes.

    python3 task307/generator/pack.py <target dir>
"""
import math
import os
import sys

import heapq
import numpy as np
import pandas as pd
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import params as P  # noqa: E402

F = P.FILES
# Table 4.2 of the model documentation: capacity per lane, alpha and beta by facility type
TABLE = {1: (1900, 0.15, 6), 12: (1900, 0.80, 6), 2: (1800, 0.80, 6), 3: (1700, 0.80, 6), 4: (850, 0.60, 4),
         5: (650, 0.75, 4), 9: (None, 0.0, 1)}
WEEKDAYS, EXPANSION = 247, 2.9          # scoring guidance 2.2


class Net:
    def __init__(self, target):
        L = pd.read_csv(os.path.join(target, F["links"]))
        nodes = pd.read_csv(os.path.join(target, F["nodes"]))
        self.zones = set(nodes.N[nodes.CENTROID == 1])
        arcs = []
        for r in L.itertuples():
            for a, b, ln in ((r.A_NODE, r.B_NODE, r.AB_LANES), (r.B_NODE, r.A_NODE, r.BA_LANES)):
                if ln > 0:
                    arcs.append(dict(link=r.LINK_ID, a=a, b=b, ftype=r.FTYPE, lanes=ln,
                                     t0=60.0 * r.LENGTH_MI / r.FF_SPD_MPH))
        self.arcs = arcs
        self.n = len(arcs)
        self.t0 = np.array([x["t0"] for x in arcs])
        self.al = np.array([TABLE[x["ftype"]][1] for x in arcs])
        self.be = np.array([TABLE[x["ftype"]][2] for x in arcs], float)
        self.cap = np.array([x["lanes"] * TABLE[x["ftype"]][0] if x["ftype"] != 9 else 1.0 for x in arcs])
        self.out = {}
        for j, x in enumerate(arcs):
            self.out.setdefault(x["a"], []).append(j)
        tr = pd.read_csv(os.path.join(target, F["trips"]))
        self.od = {per: {(int(i), int(j)): float(q) for i, j, q in zip(g.I, g.J, g.TRIPS)}
                   for per, g in tr.groupby("PERIOD")}
        wb = load_workbook(os.path.join(target, F["tip"]), read_only=True)
        cand = list(wb["Candidates"].iter_rows(values_only=True))[1:]
        self.cost = {r[0]: float(r[3]) for r in cand}
        self.names = {r[0]: r[1] for r in cand}
        self.coding = {}
        for tip, link, a, b, change, fr, to in list(wb["Model coding"].iter_rows(values_only=True))[1:]:
            self.coding.setdefault(tip, []).append((int(link), str(change), float(fr), float(to)))
        self.loaded = pd.read_csv(os.path.join(target, F["loaded"]))

    # ------------------------------------------------------------ capacities with a candidate coded in
    def caps(self, tip=None):
        c = self.cap.copy()
        if tip is None:
            return c
        for link, change, fr, to in self.coding[tip]:
            for j, x in enumerate(self.arcs):
                if x["link"] != link:
                    continue
                if change.startswith("AB_LANES"):
                    c[j] = to * TABLE[x["ftype"]][0]
                else:
                    c[j] = x["lanes"] * to
        return c

    def times(self, v, cap):
        r = np.where(self.al > 0, v / cap, 0.0)
        return self.t0 * (1 + self.al * r ** self.be)

    def deriv(self, v, cap):
        r = np.where(self.al > 0, v / cap, 0.0)
        return np.where(self.al > 0, self.t0 * self.al * self.be * r ** (self.be - 1) / cap, 0.0)

    def tree(self, o, t):
        dist = {o: 0.0}
        pred = {}
        h = [(0.0, o)]
        done = set()
        while h:
            d, u = heapq.heappop(h)
            if u in done:
                continue
            done.add(u)
            if u != o and u in self.zones:
                continue                       # a zone is an end, never a way through
            for j in self.out.get(u, []):
                w = self.arcs[j]["b"]
                nd = d + t[j]
                if nd < dist.get(w, math.inf) - 1e-15:
                    dist[w] = nd
                    pred[w] = j
                    heapq.heappush(h, (nd, w))
        return dist, pred

    @staticmethod
    def path(pred, o, d):
        p = []
        u = d
        while u != o:
            j = pred[u]
            p.append(j)
            u = None if j is None else None
            u = Net._tail[j]
        return tuple(reversed(p))

    def ue(self, od, cap, tol=1e-12, iters=4000):
        """user equilibrium by path-based gradient projection: each OD's flow moved onto its current shortest path
        from its costlier paths, a Newton step at a time, until the relative gap is below tol"""
        Net._tail = [x["a"] for x in self.arcs]
        by_o = {}
        for (o, d), q in od.items():
            if q > 0:
                by_o.setdefault(o, []).append((d, q))
        paths = {}
        v = np.zeros(self.n)
        t = self.times(v, cap)
        for o, dl in by_o.items():
            _, pred = self.tree(o, t)
            for d, q in dl:
                p = self.path(pred, o, d)
                paths[(o, d)] = {p: q}
                v[list(p)] += q
        total = sum(od.values())
        gap = 1.0
        for it in range(iters):
            t = self.times(v, cap)
            num = 0.0
            for o, dl in by_o.items():
                dist, pred = self.tree(o, t)
                for d, q in dl:
                    ps = paths[(o, d)]
                    sp = self.path(pred, o, d)
                    if sp not in ps:
                        ps[sp] = 0.0
                    for p in list(ps):
                        if p == sp or ps[p] <= 0:
                            continue
                        tt = self.times(v, cap)
                        cp = tt[list(p)].sum()
                        cs = tt[list(sp)].sum()
                        if cp - cs <= 0:
                            continue
                        dd = self.deriv(v, cap)
                        sym = set(p) ^ set(sp)
                        den = dd[list(sym)].sum()
                        step = min(ps[p], (cp - cs) / den if den > 0 else ps[p])
                        ps[p] -= step
                        ps[sp] += step
                        v[list(p)] -= step
                        v[list(sp)] += step
                    for p in [p for p, f in ps.items() if f <= 1e-12 and p != sp]:
                        del ps[p]
            t = self.times(v, cap)
            tstt = float(v @ t)
            sptt = 0.0
            for o, dl in by_o.items():
                dist, _ = self.tree(o, t)
                sptt += sum(q * dist[d] for d, q in dl)
            gap = (tstt - sptt) / tstt
            if gap < tol:
                break
        return v, self.times(v, cap), gap

    def vht(self, v, t):
        return float(v @ t) / 60.0

    def call(self, tol=1e-12):
        out = {"nobuild": {}, "proj": {}}
        nb = {}
        for per, od in self.od.items():
            v, t, g = self.ue(od, self.caps(), tol)
            nb[per] = (v, t, self.vht(v, t), g)
            out["nobuild"][per] = nb[per][2]
        for tip in self.coding:
            c = self.caps(tip)
            saved = {}
            for per, od in self.od.items():
                v, t, g = self.ue(od, c, tol)
                saved[per] = nb[per][2] - self.vht(v, t)
            annual = (saved["AM"] + saved["PM"]) * WEEKDAYS * EXPANSION
            out["proj"][tip] = dict(saved=saved, annual=annual, per_m=annual / self.cost[tip])
        out["best"] = max(out["proj"], key=lambda k: out["proj"][k]["annual"])
        return out

    def fixed(self):
        """the cheap route: each candidate's own links re-timed at the no-build loaded volumes"""
        out = {}
        for tip, rows in self.coding.items():
            s = 0.0
            for per in ("AM", "PM"):
                for link, change, fr, to in rows:
                    r = self.loaded[self.loaded.LINK_ID == link].iloc[0]
                    for j, x in enumerate(self.arcs):
                        if x["link"] != link:
                            continue
                        vol = r["AB_VOL_" + per] if x["a"] == r.A_NODE else r["BA_VOL_" + per]
                        c0 = self.cap[j]
                        c1 = to * TABLE[x["ftype"]][0] if change.startswith("AB_LANES") else x["lanes"] * to
                        f = lambda c: x["t0"] * (1 + self.al[j] * (vol / c) ** self.be[j])
                        s += vol * (f(c0) - f(c1)) / 60.0
            ann = s * WEEKDAYS * EXPANSION
            out[tip] = dict(annual=ann, per_m=ann / self.cost[tip])
        out["best"] = max(out, key=lambda k: out[k]["annual"])
        return out


if __name__ == "__main__":
    n = Net(sys.argv[1])
    c = n.call()
    print("no-build VHT", c["nobuild"])
    for k, v in c["proj"].items():
        print(k, n.names[k], "annual %.1f per_m %.1f saved %s" % (v["annual"], v["per_m"], v["saved"]))
    print("best", c["best"])
    fx = n.fixed()
    print("fixed:", {k: (round(v["annual"]), round(v["per_m"])) for k, v in fx.items() if k != "best"}, fx["best"])
