"""The three asks computed from the cut pack alone: each station's 2025 AADT and K-factor, the base-year validation at
the three river-bridge stations, and each CMP corridor's planning time index; with the paths a response takes that
does not meet each ask's device.

    python3 task307/generator/asks.py <target dir>
"""
import datetime as dt
import math
import os
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import params as P  # noqa: E402

F = P.FILES
HOLIDAYS = {dt.date(2025, 1, 1), dt.date(2025, 1, 20), dt.date(2025, 2, 17), dt.date(2025, 5, 26),
            dt.date(2025, 6, 19), dt.date(2025, 7, 4), dt.date(2025, 9, 1), dt.date(2025, 10, 13),
            dt.date(2025, 11, 11), dt.date(2025, 11, 27), dt.date(2025, 12, 25)}
BRIDGES = ["S-0101", "S-0102", "S-0103"]
VALID_MONTHS = (3, 4, 5, 9, 10)


def half_up(v, nd=0):
    q = 10 ** nd
    return math.floor(v * q + 0.5) / q


# ---------------------------------------------------------------- the count stations
class Counts:
    def __init__(self, target, df=None, inv=None):
        if df is None:
            df = pd.read_csv(os.path.join(target, F["atr"]))
        if inv is None:
            wb = load_workbook(os.path.join(target, F["stations"]), read_only=True)
            inv = {ws.title: list(ws.iter_rows(values_only=True))[1:] for ws in wb}
        rows = inv["Stations"]
        # a replacement station (suffix R) continues its predecessor's series (count manual 3.3)
        self.series = {r[0]: (r[0][:-1] if r[0].endswith("R") else r[0]) for r in rows}
        self.maint = inv["Maintenance 2025"]
        df["SER"] = df.STATION.map(self.series)
        df["D"] = pd.to_datetime(df.DATE, format="%m/%d/%Y").dt.date
        # the repeated hour on the day the clocks go back: its second occurrence
        df["OCC"] = df.groupby(["SER", "DIR", "D", "HR_END"]).cumcount()
        self.df = df
        self.stations = sorted(df.SER.unique())
        self.dirs = {s: sorted(df.loc[df.SER == s, "DIR"].unique()) for s in self.stations}
        self._wide = {}

    def wide(self, s):
        """one row per hour of the station's year: both directions' volumes and statuses side by side"""
        if s not in self._wide:
            x = self.df[self.df.SER == s]
            d1, d2 = self.dirs[s]
            a = x[x.DIR == d1].set_index(["D", "HR_END", "OCC"])
            b = x[x.DIR == d2].set_index(["D", "HR_END", "OCC"])
            w = pd.DataFrame({"v1": a.VOLUME, "s1": a.STATUS, "v2": b.VOLUME, "s2": b.STATUS}).reset_index()
            w["seq"] = np.arange(len(w))
            self._wide[s] = w
        return self._wide[s]

    # the faults the work orders closed: a loop or detector card replaced on one direction
    def fault_orders(self):
        out = []
        for st, day, wo, txt in self.maint:
            t = str(txt)
            ser = self.series.get(st, st)
            if "open circuit" in t or "crosstalk" in t:
                d = t.split()[2] if t.startswith("Loop") else t.split()[0]
                out.append((ser, d, dt.date.fromisoformat(str(day)[:10])))
        return out

    def fault_start(self, s, d, repair):
        """the first day the direction's count left its usual share of the station's two-way traffic: each hour's
        share against the median share for that hour of the week over the rest of the year, and the one
        change point in that series over the five weeks before the work order's day (the split that best separates
        the hours before it from the hours after)"""
        w = self.wide(s)
        col = "v1" if self.dirs[s][0] == d else "v2"
        w = w[(w.s1 == "OK") & (w.s2 == "OK") & ((w.v1 + w.v2) > 300)].copy()
        w["share"] = w[col] / (w.v1 + w.v2)
        w["dow"] = [x.weekday() for x in w.D]
        base = w[(w.D < repair - dt.timedelta(days=35)) | (w.D > repair)]
        ref = base.groupby(["dow", "HR_END"]).share.median()
        x = w[(w.D >= repair - dt.timedelta(days=35)) & (w.D < repair)]
        dev = np.log(x.share.values / ref.reindex(list(zip(x.dow, x.HR_END))).values)
        keep = ~np.isnan(dev)
        dev, days = dev[keep], x.D.values[keep]
        n = len(dev)
        cs = np.cumsum(dev)
        k = np.arange(1, n)
        stat = np.abs(cs[k - 1] / k - (cs[-1] - cs[k - 1]) / (n - k)) * np.sqrt(k * (n - k) / n)
        cp = int(k[np.argmax(stat)])
        return days[cp]

    def fault_windows(self):
        return {(s, d): (self.fault_start(s, d, r), r) for s, d, r in self.fault_orders() if s in self.stations}

    def hour_flags(self, s, faults=True, carry=True, carry_after_spell=False):
        """per hour: True when the hour is a true count of its traffic in both directions"""
        w = self.wide(s)
        good = (w.s1 == "OK") & (w.s2 == "OK")
        if carry:
            nr = ((w.s1 == "NR") | (w.s2 == "NR")).values
            prev = np.r_[False, nr[:-1]]
            if not carry_after_spell:
                # after a whole-day outage the counters restart: the first hour back is its own count
                run = np.zeros(len(nr), int)
                for i in range(len(nr)):
                    run[i] = run[i - 1] + 1 if nr[i] and i else int(nr[i])
                prevrun = np.r_[0, run[:-1]]
                prev = prev & (prevrun < 20)
            good &= ~prev
        if faults:
            for (ss, d), (d0, d1) in self.fault_windows().items():
                if ss == s:
                    good &= ~((w.D >= d0) & (w.D <= d1)).values
        return good.values

    def valid_days(self, s, **kw):
        w = self.wide(s)
        g = pd.Series(self.hour_flags(s, **kw)).groupby(w.D.values).all()
        return set(g[g].index)

    def daily(self, s):
        w = self.wide(s)
        return (w.v1 + w.v2).groupby(w.D.values).sum()

    def aadt(self, s, days=None, method="aashto", **kw):
        days = self.valid_days(s, **kw) if days is None else days
        dv = self.daily(s)
        dv = dv[dv.index.isin(days)]
        if method == "mean":
            return float(dv.mean())
        x = pd.DataFrame({"v": dv.values, "m": [d.month for d in dv.index], "w": [d.weekday() for d in dv.index]})
        mdw = x.groupby(["m", "w"]).v.mean()
        madt = mdw.groupby(level=0).mean()
        assert len(madt) == 12 and (mdw.groupby(level=0).size() == 7).all()
        return float(madt.mean())

    def k30(self, s, aadt, **kw):
        w = self.wide(s)
        v = (w.v1 + w.v2)[self.hour_flags(s, **kw)]
        h30 = float(np.sort(v.values)[::-1][29])
        return h30 / aadt, h30

    def ask1(self, **kw):
        out = {}
        for s in self.stations:
            a = self.aadt(s, **kw)
            k, h = self.k30(s, half_up(a), **kw)
            out[s] = dict(aadt=a, k=k, h30=h, k_raw=self.k30(s, a, **kw)[0])
        return out

    # ------------------------------------------------------------ validation at the river bridges
    def ask2(self, target, val=None, **kw):
        val = pd.read_csv(os.path.join(target, F["valid"])) if val is None else val
        out = {}
        for s in [b for b in BRIDGES if b in self.stations]:
            w = self.wide(s)
            vd = {d for d in self.valid_days(s, **kw) if d.weekday() in (1, 2, 3) and d.month in VALID_MONTHS
                  and d not in HOLIDAYS}
            for per, he in (("AM", 8), ("PM", 18)):
                x = w[(w.HR_END == he) & (w.OCC == 0) & w.D.isin(vd)]
                for j, d in enumerate(self.dirs[s]):
                    cnt = float(x["v%d" % (j + 1)].mean())
                    r = val[(val.STATION == s) & (val.DIR == d) & (val.PERIOD == per)].iloc[0]
                    model = r.VOL_PCE / (1 + r.HV_PCT / 100.0)
                    out[(s, d, per)] = dict(model=model, count=cnt, pct=100.0 * (model - cnt) / cnt, days=len(x))
        return out


# ---------------------------------------------------------------- the CMP corridors
class Probe:
    def __init__(self, target, seg=None, pr=None):
        self.seg = pd.read_csv(os.path.join(target, F["segs"]), dtype={"SEGMENT": str}) if seg is None else seg
        pr = pd.read_parquet(os.path.join(target, F["probe"])) if pr is None else pr
        pr["utc"] = pd.to_datetime(pr.measurement_tstamp, utc=True)
        self.pr = pr

    def pti(self, stamp="end", tz="America/Chicago", complete=True, rank="nearest", holidays=True):
        pr = self.pr
        loc = pr.utc.dt.tz_convert(tz).dt.tz_localize(None) if tz else pr.utc.dt.tz_localize(None)
        start = loc - pd.Timedelta(minutes=15) if stamp == "end" else loc
        x = pd.DataFrame({"SEGMENT": pr.segment_id.values, "start": start.values, "tt": pr.travel_time_seconds.values})
        x["day"] = x.start.dt.date
        x["hour"] = x.start.dt.hour
        x = x[(x.start.dt.weekday < 5)]
        if holidays:
            x = x[~x.day.isin(HOLIDAYS)]
        out = {}
        for cor in sorted(set(self.seg.CORRIDOR) & set(self.seg.loc[self.seg.SEGMENT.isin(set(pr.segment_id)), "CORRIDOR"])):
            for per, hr, inbound in (("AM", 7, True), ("PM", 17, False)):
                sg = self.seg[(self.seg.CORRIDOR == cor) & (self.seg.DIR.isin(["NB", "EB"]) == inbound)]
                ff = float((3600.0 * sg.LENGTH_MI / sg.REF_SPEED_MPH).sum())
                y = x[x.SEGMENT.isin(set(sg.SEGMENT)) & (x.hour == hr)]
                g = y.groupby("start").tt.agg(["sum", "size"])
                if complete:
                    g = g[g["size"] == len(sg)]
                tt = np.sort(g["sum"].values)
                n = len(tt)
                p95 = float(tt[math.ceil(0.95 * n) - 1]) if rank == "nearest" else float(np.percentile(tt, 95))
                out[(cor, per)] = dict(pti=p95 / ff, p95=p95, ff=ff, n=n)
        return out


def golden(target):
    c = Counts(target)
    return dict(ask1=c.ask1(), ask2=c.ask2(target), ask3=Probe(target).pti(), windows=c.fault_windows())


def show(target):
    c = Counts(target)
    print("fault windows:", {k: (str(a), str(b)) for k, (a, b) in c.fault_windows().items()})
    g1 = c.ask1()
    paths1 = {"filed rules only (the cracker)": c.ask1(faults=False, carry=False),
              "faults from the data, carried polls kept": c.ask1(carry=False),
              "carried polls dropped, faults kept": c.ask1(faults=False),
              "every hour after an NR dropped, spells too": c.ask1(carry_after_spell=True)}
    for s in c.stations:
        print("%s AADT %.2f K %.4f (raw %.4f) h30 %.0f" % (s, g1[s]["aadt"], g1[s]["k"], g1[s]["k_raw"], g1[s]["h30"]))
    for name, p in paths1.items():
        ha = sum(half_up(p[s]["aadt"]) == half_up(g1[s]["aadt"]) for s in c.stations)
        hk = sum(half_up(p[s]["k"], 3) == half_up(g1[s]["k"], 3) for s in c.stations)
        print("  ask 1 %-48s AADT %d of 8, K %d of 8" % (name, ha, hk))
    g2 = c.ask2(target)
    for k, v in g2.items():
        print("%s model %.1f count %.2f pct %.3f days %d" % (k, v["model"], v["count"], v["pct"], v["days"]))
    p2 = c.ask2(target, faults=False, carry=False)
    print("  ask 2 filed rules only: %d of 12" % sum(half_up(p2[k]["pct"], 1) == half_up(g2[k]["pct"], 1) for k in g2))
    pb = Probe(target)
    g3 = pb.pti()
    for k, v in g3.items():
        print("%s PTI %.4f n %d" % (k, v["pti"], v["n"]))
    for name, kw in {"stamp read as the start": dict(stamp="start"), "UTC read as local": dict(tz=None),
                     "interpolated percentile": dict(rank="linear"), "incomplete intervals kept": dict(complete=False)}.items():
        p = pb.pti(**kw)
        print("  ask 3 %-30s %d of 8" % (name, sum(half_up(p[k]["pti"], 2) == half_up(g3[k]["pti"], 2) for k in g3)))


if __name__ == "__main__":
    show(sys.argv[1])
