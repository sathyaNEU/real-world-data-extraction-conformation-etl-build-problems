"""The two deliverables the planning desk would put in front of the board."""
import os

import world as W


def workbook(path, g):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Augment decision"
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"] = "%s: interconnect augment decision, first quarter 2027" % W.PLATFORM
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Worst evening shortfall by interconnect and evening window, the evenings the "
                "network operations team had booked work left out.")
    ws["A2"].font = Font(size=9, italic=True)
    ws["A4"] = "Recommendation"
    ws["A4"].font = Font(bold=True)
    ws["B4"] = ("Fund no augment this cycle" if g["fund"] == "none"
                else "Fund the augment on %s" % g["fund"])
    ws["A5"] = "Comes closest"
    ws["A5"].font = Font(bold=True)
    ws["B5"] = "%s, %d kbps" % (g["closest_port"], int(g["closest_shortfall"]))
    ws["A6"] = "Evenings left out"
    ws["A6"].font = Font(bold=True)
    ws["B6"] = "%d, on which work was booked inside an evening window" % g["evenings_left_out"]
    head = ["Interconnect", "Evening window", "Worst shortfall, kbps", "Evening it fell on"]
    for j, h in enumerate(head):
        c = ws.cell(row=8, column=1 + j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8E8E4")
        c.border = box
        c.alignment = Alignment(horizontal="center" if j else "left")
    r = 9
    for (line, win), (v, day) in sorted(g["grid"].items()):
        for j, val in enumerate([line, "%s to %02d:59" % (win, int(win[:2])), v, day]):
            c = ws.cell(row=r, column=1 + j, value=val)
            c.border = box
            if j in (2, 3):
                c.alignment = Alignment(horizontal="center")
        r += 1
    ws.cell(row=r + 1, column=1,
            value="The shortfall of an interconnect in a window is the reference band median less "
                  "the window median, on the delivered quality standard's definition.").font = \
        Font(size=9, italic=True)
    ws.cell(row=r + 2, column=1,
            value="Client telemetry to 12 April 2027. Network planning desk, for the capacity "
                  "committee of 19 April.").font = Font(size=9, italic=True)
    for col, wd in (("A", 15), ("B", 20), ("C", 22), ("D", 20)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A9"
    wb.save(path)
    return path


def chart(path, g):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    worst = {}
    for (line, win), (v, day) in g["grid"].items():
        worst[line] = max(worst.get(line, 0), v)
    order = sorted(worst, key=lambda k: -worst[k])
    vals = [worst[k] for k in order]
    fig, ax = plt.subplots(figsize=(8.4, 4.4), dpi=150)
    cols = ["#1f4e79" if k == g["closest_port"] else "#9fb6cd" for k in order]
    ax.bar(order, vals, color=cols, edgecolor="#35485c", linewidth=0.6)
    ax.axhline(int(g["record_shortfall"]), color="#8c2f2f", linestyle="--", linewidth=1.2)
    ax.annotate("the largest the record shows, %d kbps" % int(g["record_shortfall"]),
                xy=(len(order) - 1, int(g["record_shortfall"])), xytext=(-6, 6),
                textcoords="offset points", ha="right", fontsize=8.5, color="#8c2f2f")
    for i, v in enumerate(vals):
        ax.annotate("%d" % v, xy=(i, v), xytext=(0, 3), textcoords="offset points",
                    ha="center", fontsize=8.5)
    ax.set_ylabel("Worst evening shortfall, kbps")
    ax.set_title("Worst evening shortfall by interconnect, first quarter 2027")
    fig.text(0.012, 0.012, "Client telemetry, evenings with booked work left out. The dashed line "
                           "is the largest any quarter in the record produced at the same reach.",
             fontsize=7.6, color="#55606b")
    ax.set_ylim(0, max(vals + [int(g["record_shortfall"])]) * 1.16)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(axis="x", labelsize=9)
    fig.tight_layout(rect=(0, 0.035, 1, 1))
    fig.savefig(path, metadata={"Software": None, "Creation Time": None})
    plt.close(fig)
    return path
