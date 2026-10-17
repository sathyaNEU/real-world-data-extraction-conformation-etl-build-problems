"""The folder the environmental engineer was handed: the meters' exports, the survey paperwork, the plant's
own records, the met mast, and two files the question does not use."""
import csv as _csv
import datetime as dt
import math
import os
import random

import numpy as np

import sim as SM
import world as W

FILES = {
    "levels": "sound_levels_1min_2026.parquet",
    "deployments": "meter_deployments_2026.csv",
    "positions": "monitoring_positions.xlsx",
    "plan": "noise_management_plan.pdf",
    "notes": "survey_notes.md",
    "events": "plant_equipment_events_2026.csv",
    "scada": "scada_15min_2026.csv",
    "weather": "weather_10min_2026.csv",
    "bands": "third_octave_15min_2026.csv",
    "calibration": "calibration_records_2026.csv",
    "complaints": "complaints_register_2026.xlsx",
    "effluent": "effluent_quality_2026.csv",
    "manifest": "export_manifest.csv",
}
BANDS = [25, 31.5, 40, 50, 63, 80, 100, 125, 160, 200, 250, 315, 400, 500, 630, 800, 1000, 1250,
         1600, 2000, 2500, 3150, 4000, 5000, 6300, 8000, 10000, 12500]


def _write_csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = _csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)
    return path


def utc(stamp_local):
    return stamp_local - dt.timedelta(hours=SM.LOCAL_OFFSET)


def write_levels(target, series):
    """The meters' own export: one row a minute, stamped UTC, as the download software writes it."""
    import pyarrow as pa
    import pyarrow.parquet as pq
    meters, stamps, laeq, la90, lamax, lamin = [], [], [], [], [], []
    dup = None
    for sid in sorted(series):
        s = series[sid]
        lg = SM.logged(s)
        n = len(s["stamps"])
        meters.extend([W_serial(s["dep"]["meter"])] * n)
        stamps.extend([utc(x) for x in s["stamps"]])
        laeq.extend(lg["laeq"].tolist())
        la90.extend(lg["la90"].tolist())
        lamax.extend(lg["lamax"].tolist())
        lamin.extend(lg["lamin"].tolist())
        if (s["dep"]["code"], s["dep"]["k"]) == W.DUP_DEPLOYMENT:
            m = np.array([(x.date() == W.DUP_DAY) and (W.DUP_FROM <= x.hour < W.DUP_TO)
                          for x in s["stamps"]])
            dup = {"meter": W_serial(s["dep"]["meter"]),
                   "stamps": [utc(x) for x, k in zip(s["stamps"], m) if k],
                   "laeq": lg["laeq"][m].tolist(), "la90": lg["la90"][m].tolist(),
                   "lamax": lg["lamax"][m].tolist(), "lamin": lg["lamin"][m].tolist()}
    if dup:
        meters.extend([dup["meter"]] * len(dup["stamps"]))
        stamps.extend(dup["stamps"])
        laeq.extend(dup["laeq"]); la90.extend(dup["la90"])
        lamax.extend(dup["lamax"]); lamin.extend(dup["lamin"])
    order = sorted(range(len(stamps)), key=lambda i: (meters[i], stamps[i]))
    tbl = pa.table({
        "meter_serial": pa.array([meters[i] for i in order], pa.string()),
        "measured_at_utc": pa.array([stamps[i] for i in order], pa.timestamp("s")),
        "laeq_db": pa.array([laeq[i] for i in order], pa.float32()),
        "la90_db": pa.array([la90[i] for i in order], pa.float32()),
        "lamax_db": pa.array([lamax[i] for i in order], pa.float32()),
        "lamin_db": pa.array([lamin[i] for i in order], pa.float32()),
    })
    path = os.path.join(target, FILES["levels"])
    pq.write_table(tbl, path, compression="snappy", version="2.6")
    return path, len(order), dup


def W_serial(meter_id):
    return [m["serial"] for m in W.METERS if m["id"] == meter_id][0]


def write_deployments(target, deps):
    rows = []
    for d in sorted(deps, key=lambda x: x["id"]):
        m = [x for x in W.METERS if x["id"] == d["meter"]][0]
        rem = []
        if (d["code"], d["k"]) == W.DUP_DEPLOYMENT:
            rem.append("card read error on collection, download repeated")
        if d["code"] == "NSL08" and d["k"] == 0:
            rem.append("mic 3 m west of the scheduled spot and at 1.5 m, stock had reached the tripod")
        if d["code"] == "NSL04" and d["k"] == 1:
            rem.append("windshield replaced at mid-week check")
        if d["code"] == "NSL13" and d["k"] == 1:
            rem.append("tripod dropped to 1.2 m, garden wall rebuilt")
        for a, mins, why in d["nodata"]:
            rem.append("%s, %s to %s" % (why, a.strftime("%d/%m %H:%M"),
                                         (a + dt.timedelta(minutes=mins)).strftime("%H:%M")))
        rows.append([d["id"], m["serial"], d["code"], d["start_dt"].strftime("%Y-%m-%d %H:%M"),
                     d["end_dt"].strftime("%Y-%m-%d %H:%M"),
                     "%.1f" % d["mount"], "%.1f" % d["cal_before"], "%.1f" % d["cal_after"],
                     W.PEOPLE[d["crew"]][1], "; ".join(rem)])
    return _write_csv(os.path.join(target, FILES["deployments"]),
                      ["deployment_ref", "meter_serial", "position_code", "deployed_at",
                       "collected_at", "mic_height_m", "cal_check_before_db", "cal_check_after_db",
                       "operator", "remarks"], rows)


def write_positions(target):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Monitoring positions"
    ws["A1"] = "%s %s: noise monitoring positions" % (W.OPERATOR, W.SITE)
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Schedule held under the site noise management plan, clause 2.1. Grid references are Irish " \
               "Transverse Mercator. Distances are to the blower house gable."
    ws["A3"] = "Maintained by the environmental engineer. Positions with a complaint reference were added " \
               "after the reference was logged. Mic height is the height set for the position; a " \
               "deployment that stood at another height says so in the deployment record."
    head = ["position_code", "dwelling_ref", "dwelling", "easting", "northing", "distance_m",
            "bearing_deg", "mic_height_m", "mounting", "screening", "complaint_ref"]
    for j, h in enumerate(head):
        c = ws.cell(row=5, column=1 + j, value=h)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="left")
    rng = random.Random(W.SEED + 5)
    for i, p in enumerate(W.POSITIONS):
        code, dref, desc, dist, bear = p[0], p[1], p[2], p[3], p[4]
        # the grid reference is where the tripod stood, surveyed on the ground; the distance and bearing
        # are to the blower house gable, scaled off the site plan to the nearest five metres and degree,
        # so the two do not reduce to one another
        ang = math.radians(bear + rng.uniform(-1.4, 1.4))
        d = dist + rng.uniform(-4.0, 4.0)
        e = 612430 + int(round(d * math.sin(ang)))
        n = 748215 + int(round(d * math.cos(ang)))
        cref = None if W.COMPLAINTS[code] == 0 else "C%02d" % (1 + sorted(
            [c for c in W.CODES if W.COMPLAINTS[c] > 0], key=lambda c: -W.COMPLAINTS[c]).index(code))
        r = 6 + i
        for j, v in enumerate([code, dref, desc, e, n, dist, bear, W.HEIGHT[code],
                               "free field, 3.5 m from facade", W.SCREEN[code], cref]):
            ws.cell(row=r, column=1 + j, value=v)
    for col, wd in zip("ABCDEFGHIJK", (14, 12, 34, 9, 9, 11, 12, 13, 28, 42, 13)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A6"
    path = os.path.join(target, FILES["positions"])
    wb.save(path)
    return path


def write_events(target):
    """The process operator's equipment log, local time, as the shift writes it. The weekly inspection
    stops the blowers and the engine; the centrifuges are already down at that hour and come back behind
    the inspection."""
    rows = []
    day = dt.date(2026, 8, 24)
    end = dt.date(2026, 10, 16)
    while day <= end:
        runs = SM.centrifuge_runs(day)
        for i, (cf, a, b) in enumerate(runs):
            same = [r for r in runs if r[0] == cf]
            second = len(same) > 1 and same.index((cf, a, b)) == 1
            why = ("dewatering restarted, liquor tank drawn down" if second else
                   "dewatering shift start behind the weekly inspection" if a.hour >= 7 else
                   "dewatering shift start")
            rows.append([a, cf, "START", W.PEOPLE["operator"][1], why])
            rows.append([b, cf, "STOP", W.PEOPLE["operator"][1],
                         "stopped to draw the liquor tank down" if (len(same) > 1 and not second)
                         else "dewatering shift end"])
        day += dt.timedelta(days=1)
    for start, mins, reason in W.PLANNED_STOPS + W.UNPLANNED_STOPS:
        planned = any(start == s for s, _, _ in W.PLANNED_STOPS)
        who = W.PEOPLE["planner"][1] if planned else W.PEOPLE["operator"][1]
        for eq in W.BLOWERS + [W.ENGINE]:
            rows.append([start, eq, "STOP", who, reason])
            rows.append([start + dt.timedelta(minutes=mins), eq, "START", who,
                         "returned to service after %s" % reason])
    for start, mins, reason in W.BL12_STOPS:
        for eq in ("BL-01", "BL-02"):
            rows.append([start, eq, "STOP", W.PEOPLE["operator"][1], reason])
            rows.append([start + dt.timedelta(minutes=mins), eq, "START", W.PEOPLE["operator"][1],
                         "returned to service after %s" % reason.split(", ")[1]])
    for start, mins, reason in W.ENGINE_STOPS:
        rows.append([start, W.ENGINE, "STOP", W.PEOPLE["operator"][1], reason])
        rows.append([start + dt.timedelta(minutes=mins), W.ENGINE, "START", W.PEOPLE["operator"][1],
                     "returned to service after %s" % reason.split(", ")[0]])
    for eq in W.BLOWERS + [W.ENGINE]:
        rows.append([dt.datetime(2026, 8, 24, 6, 12), eq, "START", W.PEOPLE["operator"][1],
                     "returned to service after oil change"])
    rows.sort(key=lambda r: (r[0], r[1], r[2]))
    out = [[r[0].strftime("%Y-%m-%d %H:%M"), r[1], r[2], r[3], r[4]] for r in rows]
    return _write_csv(os.path.join(target, FILES["events"]),
                      ["logged_at", "equipment", "event", "logged_by", "reason"], out)


def _ar1(n, a, rng):
    """Unit-variance AR(1) noise: a process channel wanders, it does not jump about independently every
    quarter of an hour."""
    out = np.zeros(n)
    wn = rng.normal(0.0, math.sqrt(1 - a * a), n)
    for i in range(n):
        out[i] = a * out[i - 1] + wn[i] if i else wn[i]
    return out


def write_scada(target):
    """Fifteen-minute process channels off the plant PLC.

    The aeration air is metered twice: the header meter that was there before the upgrade, and the third
    blower's own meter, which was fitted when BL-03 was tied into the main downstream of it. The hour
    counters on the gas skid were tagged from the contractor's own I/O list. The biogas meter is a counter
    and rolls over, one hour counter was replaced mid-survey, and the radio link to the PLC dropped
    through the survey, sometimes writing nothing and sometimes writing the last value it had.
    """
    rng = np.random.default_rng(W.SEED + 23 + W.AIR_SEED)
    t0 = dt.datetime(2026, 8, 1, 0, 0)
    n = int((dt.datetime(2026, 10, 17) - t0).total_seconds() // 900)
    stamps = [t0 + dt.timedelta(minutes=15 * i) for i in range(n)]
    stopped = np.zeros(n, dtype=bool)
    for start, e in SM.all_down():
        stopped |= np.array([(s >= start - dt.timedelta(minutes=14)) and (s < e) for s in stamps])
    hours = np.array([s.hour + s.minute / 60.0 for s in stamps])
    dows = np.array([s.weekday() for s in stamps])
    # the site's weekly milk intake moves the plant's load from week to week; the sound carries it too
    wkf = np.array([SM.week_factor(s.date()) for s in stamps])
    inflow = (96 + 34 * np.exp(-((hours - 10.5) ** 2) / 9.0) + 26 * np.exp(-((hours - 19.0) ** 2) / 7.0)
              + 6.1 * _ar1(n, 0.74, rng))
    inflow = np.where(stopped, 0.0, inflow)        # the inlet pumps are on the MCC the inspection stops
    # the air the lanes take: the dissolved oxygen trim follows the load on the works, so the total
    # wanders through the day with the inlet and over the weeks with the intake
    total_air = ((4120 + 180 * np.sin(2 * math.pi * hours / 24.0 - 1.2)
                  + 232 * _ar1(n, 0.88, rng) + 78 * _ar1(n, 0.42, rng)) * wkf
                 + 7.4 * (inflow - float(inflow.mean())))
    total_air = np.where(stopped, 0.0, np.clip(total_air, 0.0, None))
    # BL-03 holds the output its drive is set to; the header's two machines take the trim, so the two
    # meters' readings are not fractions of one another
    setting = np.array([max(v for d0, v in W.BL03_SETTING if d0 <= s.date()) for s in stamps])
    bl03 = np.where(stopped, 0.0, setting * (1.0 + 0.018 * _ar1(n, 0.8, rng)))
    air = np.clip(total_air - bl03, 0.0, None)
    # the spells the shift isolates the header's own two machines in: BL-03 carries the lanes alone, so
    # the header's meter reads nothing and its totaliser does not advance while BL-03's does
    iso = np.zeros(n, dtype=bool)
    for start, mins, _ in W.BL12_STOPS:
        end = start + dt.timedelta(minutes=mins)
        iso |= np.array([(s > start) and (s <= end) for s in stamps])
    air = np.where(iso, 0.0, air)
    # the two flow totalisers, thousands of normal cubic metres: the header's has run since the works was
    # built, BL-03's since it was commissioned at the upgrade
    air_total = 128450.0 + np.cumsum(air * 0.25) / 1000.0
    bl03_total = 2184.0 + np.cumsum(bl03 * 0.25) / 1000.0
    cf_on = np.zeros(n, dtype=bool)
    runs = {d: SM.centrifuge_runs(d) for d in sorted({s.date() for s in stamps})}
    for i, st in enumerate(stamps):
        for _, a, b in runs[st.date()]:
            if a <= st < b:
                cf_on[i] = True
                break
        cf_on[i] = cf_on[i] and not stopped[i]
    feed = np.where(cf_on, 31.5 + rng.normal(0, 1.4, n), 0.0)
    gas_rate = np.where(stopped, 0.0, (212 + 9.5 * np.sin(2 * math.pi * hours / 24.0 + 0.4)
                                       + 5.2 * _ar1(n, 0.84, rng)) * wkf)   # m3 an hour to the engine
    counter = (18640.0 + np.cumsum(gas_rate * 0.25)) % 100000.0
    eng_off = stopped.copy()
    for start, mins, _ in W.ENGINE_STOPS:
        end = start + dt.timedelta(minutes=mins)
        eng_off |= np.array([(s >= start - dt.timedelta(minutes=14)) and (s < end) for s in stamps])
    kw = np.where(eng_off, 0.0,
                  268 + 7.0 * np.sin(2 * math.pi * hours / 24.0 + 0.4) + 4.2 * _ar1(n, 0.8, rng))
    # the two hour counters, each counting the minutes its own machine turned in the quarter hour it
    # closes: the engine, which is down for the weekly inspection and for its own trips, and the gas
    # booster, which is down for the inspection and for its daily drain and purge and runs on through
    # every stop the engine takes on its own
    eng_min = _run_minutes(stamps, SM.all_down()
                           + [(s, s + dt.timedelta(minutes=m)) for s, m, _ in W.ENGINE_STOPS])
    purges = []
    day = t0.date()
    while day <= dt.date(2026, 10, 17):
        h, m = W.BOOSTER_PURGE_AT[day.weekday()]
        a = dt.datetime(day.year, day.month, day.day, h, m)
        purges.append((a, a + dt.timedelta(minutes=W.BOOSTER_PURGE_MIN)))
        day += dt.timedelta(days=1)
    bst_min = _run_minutes(stamps, SM.all_down() + purges)
    eng_hours = np.cumsum(eng_min) / 60.0
    bst_hours = np.cumsum(bst_min) / 60.0
    replaced = np.array([s >= W.HOUR_COUNTER_SWAP for s in stamps])
    base = eng_hours[np.argmax(replaced)] if replaced.any() else 0.0
    eng_hours = np.where(replaced, eng_hours - base, 41820.0 + eng_hours)
    bst_hours = 52960.0 + bst_hours
    gap = np.zeros(n, dtype=bool)
    for start, blocks in W.PLC_GAPS:
        end = start + dt.timedelta(minutes=15 * blocks)
        gap |= np.array([(s >= start) and (s < end) for s in stamps])
    held = np.zeros(n, dtype=bool)
    for start, blocks in W.PLC_HELD:
        end = start + dt.timedelta(minutes=15 * blocks)
        held |= np.array([(s >= start) and (s < end) for s in stamps])
    rows, last = [], None
    for i, s in enumerate(stamps):
        if gap[i]:
            continue
        if held[i] and last is not None:
            rows.append([s.strftime("%Y-%m-%d %H:%M")] + last)
            continue
        vals = ["%.0f" % max(air[i], 0.0), "%.1f" % air_total[i],
                "%.0f" % max(bl03[i], 0.0), "%.1f" % bl03_total[i],
                "%.1f" % max(feed[i], 0.0), "%.0f" % counter[i], "%.1f" % max(kw[i], 0.0),
                "%.1f" % eng_hours[i], "%.1f" % bst_hours[i], "%.1f" % max(inflow[i], 0.0)]
        rows.append([s.strftime("%Y-%m-%d %H:%M")] + vals)
        last = vals
    return _write_csv(os.path.join(target, FILES["scada"]),
                      ["logged_at", "aeration_header_nm3_h", "aeration_header_total_knm3",
                       "bl03_discharge_nm3_h", "bl03_discharge_total_knm3",
                       "centrifuge_feed_m3_h", "biogas_meter_m3", "engine_export_kw",
                       "engine_hours_run", "gas_booster_hours_run", "inlet_flow_m3_h"], rows)


def _run_minutes(stamps, down):
    """Minutes of the quarter hour ending at each stamp that the machine was turning, given the intervals
    it was down for."""
    out = np.zeros(len(stamps))
    for i, s in enumerate(stamps):
        a = s - dt.timedelta(minutes=15)
        mins = 15.0
        for lo, hi in down:
            lo2, hi2 = max(a, lo), min(s, hi)
            if hi2 > lo2:
                mins -= (hi2 - lo2).total_seconds() / 60.0
        out[i] = max(mins, 0.0)
    return out


def write_weather(target):
    """The met mast, UTC, ten-minute means. The anemometer was swapped and reports in km an hour after it;
    the rain gauge reports its tip count since midnight."""
    rows = []
    for i in range(SM.MET["n"]):
        stamp = SM.MET["t0"] + dt.timedelta(minutes=10 * i)
        rows.append([stamp.strftime("%Y-%m-%d %H:%M"), "%.1f" % SM.MET["wind_shipped"][i],
                     "%d" % SM.MET["wind_dir"][i],
                     "%d" % SM.MET["tips"][i], "%.1f" % SM.MET["temp_c"][i]])
    return _write_csv(os.path.join(target, FILES["weather"]),
                      ["measured_at_utc", "wind_speed", "wind_direction_deg", "rain_tips_today",
                       "air_temperature_c"], rows)


A_WEIGHT = {25: -44.7, 31.5: -39.4, 40: -34.6, 50: -30.2, 63: -26.2, 80: -22.5, 100: -19.1,
            125: -16.1, 160: -13.4, 200: -10.9, 250: -8.6, 315: -6.6, 400: -4.8, 500: -3.2,
            630: -1.9, 800: -0.8, 1000: 0.0, 1250: 0.6, 1600: 1.0, 2000: 1.2, 2500: 1.3,
            3150: 1.2, 4000: 1.0, 5000: 0.5, 6300: -0.1, 8000: -1.1, 10000: -2.5, 12500: -4.3}


def _spectra(code, wobble):
    """Unweighted band energy fractions for the two sources at a position: the plant, with the blower's
    blade tone at 400 Hz and the engine's firing tone at 100 Hz as far as they carry to the position, and
    the surroundings, which are road traffic and wind in trees and carry no tone."""
    k400, k100 = W.TONE_K[code]
    pl, rs = [], []
    for i, b in enumerate(BANDS):
        f = float(b)
        p = -8.5 - 11.0 * abs(math.log10(f / 315.0)) ** 1.85 + wobble[i]
        if abs(f - 400.0) < 1.0:
            p += 16.0 * k400
        elif abs(f - 100.0) < 1.0:
            p += 5.9 * k100
        elif abs(f - 800.0) < 1.0:
            p += 4.0 * k400 * k400
        r = -7.0 - 8.0 * abs(math.log10(f / 100.0)) ** 1.62 + 0.6 * wobble[i]
        pl.append(10.0 ** (p / 10.0))
        rs.append(10.0 ** (r / 10.0))
    pl = np.array(pl) / sum(pl)
    rs = np.array(rs) / sum(rs)
    aw = np.array([10.0 ** (A_WEIGHT[b] / 10.0) for b in BANDS])
    return pl, rs, float((pl * aw).sum()), float((rs * aw).sum())


def write_bands(target, series):
    """One-third-octave band levels, quarter hour, unweighted, on the three meters that carry the module."""
    rows = []
    licensed = {m["id"] for m in W.METERS if m["bands"]}
    for sid in sorted(series):
        s = series[sid]
        if s["dep"]["meter"] not in licensed:
            continue
        dep = s["dep"]
        rng = np.random.default_rng(W.SEED + 29 + 613 * int(dep["id"][1:]))
        # the mic's own position in the garden moves the spectrum a little week to week
        wob = rng.normal(0.0, 0.42, len(BANDS))
        pl, rs, apl, ars = _spectra(dep["code"], wob)
        gain = 10.0 ** (SM.dep_dev(dep) / 10.0)
        stamps = s["stamps"]
        pe = s["plant_e"] * gain
        re = s["resid_e"] * gain
        ce = s["cal_e"] * gain
        he = s["hand_e"] * gain
        nd = s["nd"]
        i1k = BANDS.index(1000)
        # the module logs on the clock, so its first quarter hour is the first one that starts on a
        # quarter of an hour after the meter was set running
        q0 = (-stamps[0].minute) % 15
        for q in range(q0, len(stamps) - 14, 15):
            sl = slice(q, q + 15)
            if nd[sl].any():
                continue
            ep = float(np.mean(pe[sl])) / apl
            er = float(np.mean(re[sl])) / ars
            vals = ep * pl + er * rs
            # the field check and the handling: a 1 kHz tone on the microphone, and the tripod itself
            ec = float(np.mean(ce[sl]))
            eh = float(np.mean(he[sl]))
            if ec > 0.0:
                vals = vals.copy()
                vals[i1k] += ec
            if eh > 0.0:
                vals = vals + (eh / ars) * rs
            vals = 10.0 * np.log10(vals) + rng.normal(0.0, 0.3, len(BANDS))
            rows.append([W_serial(dep["meter"]), utc(stamps[q]).strftime("%Y-%m-%d %H:%M")]
                        + ["%.1f" % v for v in vals])
    rows.sort(key=lambda r: (r[0], r[1]))
    return _write_csv(os.path.join(target, FILES["bands"]),
                      ["meter_serial", "measured_at_utc"] + ["lzeq_%s_hz" % str(b).replace(".", "p")
                                                             for b in BANDS], rows)


def write_calibration(target):
    """The laboratory certificates for the five meters and the field calibrator. The field checks either
    side of a deployment are on the deployment record, not here."""
    rows = []
    for m in W.METERS:
        rows.append([m["serial"], m["model"], m["cert"], m["due"], "CAL-4621",
                     "%.1f" % W.CAL_REF, "%.1f" % m["found"], "%.1f" % m["left"],
                     "periodic verification to IEC 61672-3, class 1"])
    rows.append(["CAL-4621", "Class 1 acoustic calibrator, 1 kHz", "2026-02-19", "2028-02-18", "",
                 "%.1f" % W.CAL_REF, "", "", "calibrator output verified, 94.0 and 114.0 dB"])
    return _write_csv(os.path.join(target, FILES["calibration"]),
                      ["instrument_serial", "model", "verified_on", "verification_due",
                       "calibrator_serial", "calibrator_level_db", "as_found_db", "as_left_db",
                       "note"], rows)


# The order the register's own contacts are written in: every contact is somebody's own words, so no two
# rows carry the same sentence.
TEXT_ORDER = [5, 17, 0, 9, 21, 13, 2, 19, 8, 11, 23, 3, 15, 6, 20, 1, 12, 22, 7, 16, 4, 18, 10, 14]


def write_complaints(target):
    """The complaints register. The question does not use it (metadata.json declares it)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Complaints 2026"
    ws["A1"] = "%s %s: noise complaints received" % (W.OPERATOR, W.SITE)
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Logged under the site noise management plan, clause 6.1. One row per contact."
    head = ["complaint_ref", "received_on", "dwelling_ref", "received_via", "received_by",
            "description", "response_sent_on"]
    for j, h in enumerate(head):
        ws.cell(row=4, column=1 + j, value=h).font = Font(bold=True)
    texts = [
        "Hum carrying across the yard again from about nine last night.",
        "Low drone, worse when the wind comes up from the plant. Noticed it most of the weekend.",
        "Caller cannot open the bedroom window at night for the noise.",
        "Whine off the new building, started after the upgrade went in.",
        "Humming all day Tuesday and Wednesday. Caller says it has got worse since July.",
        "Noise at the gable end of the house, caller asked what the plant is running at night.",
        "Drone in the evening. Two children not sleeping.",
        "Steady hum from the direction of the works, heard from the back door.",
        "Caller rang to ask whether the plant is running differently. Noise louder than last year.",
        "Rumble carrying down the lane, worst between six and eight in the morning.",
        "Cannot sit out in the garden with the hum, caller has given up on the patio.",
        "A tone off the plant, there all weekend and still there this morning.",
        "Hum through the floor of the front room at night.",
        "Noise from the works louder in the mornings this week.",
        "Whining noise, caller says it is getting worse month by month.",
        "The drone starts before six and does not stop.",
        "Steady noise on still nights. Caller asked for the plant to be turned down after ten.",
        "Humming from the plant heard all week, caller wants to know what is being done.",
        "Noise carries up to the house on a calm day, worse since the dry spell.",
        "A hum audible with the windows shut.",
        "Constant drone from the direction of the plant since Saturday.",
        "Noise off the works at the back of the house, caller rang twice.",
        "A whine from the new plant, caller says it was not there last autumn.",
        "Humming noise. Caller and her husband are not able to sleep with it.",
    ]
    channels = ["telephone", "telephone", "email", "telephone", "letter", "email", "telephone",
                "telephone", "email", "telephone", "telephone", "email", "telephone", "letter",
                "telephone", "email", "telephone", "telephone", "email", "telephone", "telephone",
                "email", "telephone", "letter"]
    takers = ["engineer", "engineer", "operations", "engineer", "operations", "engineer",
              "operations", "engineer", "engineer", "operations", "engineer", "engineer",
              "operations", "engineer", "engineer", "operations", "engineer", "engineer",
              "operations", "engineer", "operations", "engineer", "engineer", "operations"]
    rng = random.Random(W.SEED + 31)
    rows, n = [], 0
    ordered = sorted([c for c in W.CODES if W.COMPLAINTS[c] > 0], key=lambda c: -W.COMPLAINTS[c])
    total = sum(W.COMPLAINTS[c] for c in ordered)
    span = (dt.date(2026, 9, 28) - dt.date(2026, 6, 4)).days
    for code in ordered:
        ref = "C%02d" % (1 + ordered.index(code))
        for k in range(W.COMPLAINTS[code]):
            d = dt.date(2026, 6, 4) + dt.timedelta(days=int(round(span * n / float(total))) + k)
            sent = min(d + dt.timedelta(days=rng.choice([4, 6, 9])), dt.date(2026, 10, 9))
            j = TEXT_ORDER[n % len(TEXT_ORDER)]
            rows.append([ref, d.isoformat(), W.BY_CODE[code][1], channels[j],
                         W.PEOPLE[takers[j]][1], texts[j], sent.isoformat()])
            n += 1
    rows.sort(key=lambda r: (r[1], r[0]))
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            ws.cell(row=5 + i, column=1 + j, value=v)
    for col, wd in zip("ABCDEFG", (14, 13, 13, 13, 12, 76, 17)):
        ws.column_dimensions[col].width = wd
    path = os.path.join(target, FILES["complaints"])
    wb.save(path)
    return path


def write_effluent(target):
    """The works' discharge monitoring, as the laboratory reports it. Nothing in the question reads it."""
    rng = np.random.default_rng(W.SEED + 37)
    LOD = {"bod": 2.0, "cod": 5.0, "ss": 2.0, "amm": 0.05}
    rows = []
    day = dt.date(2026, 7, 1)
    seq = 0
    while day <= dt.date(2026, 10, 15):
        if day.weekday() in (0, 2, 4):
            seq += 1
            grab = (seq % 7 == 3)
            pt = "final effluent" if not grab else rng.choice(["final effluent", "balance tank outfall"])
            bod = max(0.4, rng.normal(7.4 if grab else 9.8, 2.6))
            cod = bod * rng.uniform(2.4, 4.1) + max(0.0, rng.normal(6.0, 2.4))
            ss = max(0.3, rng.normal(3.1, 0.9))
            amm = max(0.01, rng.normal(0.42, 0.19))
            def val(x, lod, fmt):
                return ("<%s" % (fmt % lod)) if x < lod else (fmt % x)
            note = ""
            if seq % 11 == 5:
                note = "sample received outside the 24 h holding time, result indicative"
            elif seq % 13 == 8:
                note = "suspended solids filter blinded, result from duplicate"
            rows.append([day.isoformat(), pt, "grab" if grab else "composite, 24 hour",
                         "Nore Valley Laboratories", "NV26-%05d" % (41200 + 37 * seq),
                         val(bod, LOD["bod"], "%.1f"), val(cod, LOD["cod"], "%.0f"),
                         val(ss, LOD["ss"], "%.1f"), val(amm, LOD["amm"], "%.2f"),
                         "%.1f" % rng.normal(7.3, 0.2), note])
        day += dt.timedelta(days=1)
    return _write_csv(os.path.join(target, FILES["effluent"]),
                      ["sampled_on", "sample_point", "sample_type", "laboratory", "lab_reference",
                       "bod_mg_l", "cod_mg_l", "suspended_solids_mg_l", "ammonia_n_mg_l", "ph",
                       "note"], rows)


def write_manifest(target, counts):
    """What was copied into the folder handed to the environmental engineer, where each file came from and
    on what terms. Site systems and survey paperwork are the operator's own records."""
    OWN = "Tulcarra Foods internal, not for release outside the site"
    rows = [
        [FILES["levels"], "parquet", counts["levels"], "31/08/2026 to 12/10/2026",
         "meter download software, five instruments", "2026-10-13", OWN,
         "every logged minute from the five meters, stamps in UTC"],
        [FILES["deployments"], "csv", counts["deployments"], "31/08/2026 to 12/10/2026",
         "survey field sheets, transcribed", "2026-10-13", OWN,
         "which meter stood at which position, with its field checks"],
        [FILES["positions"], "xlsx", len(W.POSITIONS), "current",
         "monitoring position schedule, environmental engineer", "2026-10-14", OWN,
         "the schedule held under the plan, clause 2.1"],
        [FILES["plan"], "pdf", "", "effective 01/06/2026",
         "site document register, revision 4", "2026-10-14", OWN,
         "the site noise management plan, clauses 1 to 7"],
        [FILES["notes"], "md", "", "31/08/2026 to 11/10/2026",
         "monitoring technician", "2026-10-16", OWN, "the technician's notes on the survey"],
        [FILES["events"], "csv", counts["events"], "24/08/2026 to 16/10/2026",
         "process operator equipment log, local time", "2026-10-16", OWN,
         "a row for every start and stop of a blower, a centrifuge or the gas engine"],
        [FILES["scada"], "csv", counts["scada"], "01/08/2026 to 16/10/2026",
         "plant PLC historian, local time", "2026-10-16", OWN,
         "fifteen-minute process channels, counters as the meters read them"],
        [FILES["weather"], "csv", counts["weather"], "30/08/2026 to 16/10/2026",
         "met mast on the blower house roof, UTC", "2026-10-16", OWN,
         "ten-minute means, wind, direction, rain gauge count and air temperature"],
        [FILES["bands"], "csv", counts["bands"], "31/08/2026 to 12/10/2026",
         "meter download software, band module", "2026-10-13", OWN,
         "one-third-octave band levels from the three meters that carry the module"],
        [FILES["calibration"], "csv", len(W.METERS) + 1, "2025-11-27 to 2026-04-08",
         "instrument certificates, document register", "2026-10-14", OWN,
         "laboratory verification of the five meters and the field calibrator"],
        [FILES["complaints"], "xlsx", counts["complaints"], "01/01/2026 to 28/09/2026",
         "complaints register, environmental engineer", "2026-10-14", OWN,
         "noise complaints received in 2026, one row a contact"],
        [FILES["effluent"], "csv", counts["effluent"], "01/07/2026 to 15/10/2026",
         "Nore Valley Laboratories reports, discharge monitoring", "2026-10-14",
         "Tulcarra Foods internal; laboratory reports used under the testing contract",
         "final effluent and balance tank results for the licence"],
    ]
    return _write_csv(os.path.join(target, FILES["manifest"]),
                      ["file", "format", "rows", "covers", "source", "extracted_on", "licence",
                       "contents"], rows)


def set_mtimes(target):
    stamp = dt.datetime(2026, 10, 16, 17, 25, 0).timestamp()
    for name in sorted(os.listdir(target)):
        os.utime(os.path.join(target, name), (stamp, stamp))
