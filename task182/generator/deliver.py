"""Write the golden deliverables into task182/golden/ from the shipped bundle alone: the bore assessment workbook
(bore_impact_assessment.xlsx, one row per register bore) and the note to the general manager
(make_good_recommendation.docx)."""
import datetime as dt
import math
import os
import sys

import openpyxl
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import analysis as A       # noqa: E402
import golden as G         # noqa: E402
import postprocess as PP   # noqa: E402

TASK = os.path.dirname(HERE)
WORKBOOK = "bore_impact_assessment.xlsx"
MEMO = "make_good_recommendation.docx"
AUTHOR = "Rebecca Quinn"
UNIT_WORD = {"alluvium": "alluvium", "basalt": "basalt", "coal": "coal measures"}
OFFER_REF = "MG-26/%02d"


def r2(x):
    return float("%.2f" % x)


def r1(x):
    return float("%.1f" % x)


def figures(g):
    """Every figure the deliverables print, from the solve, at its printed rounding."""
    obs, bores = g["obs"], g["bores"]
    rows = []
    offer_no = {}
    for owner in sorted(g["offers"], key=lambda o: min(g["offers"][o])):
        offer_no[owner] = len(offer_no) + 1
    for k in g["LH"]:
        b = bores[k]
        o = obs[k]
        sep = o["logger"] is not None
        if sep:
            basis = "logger daily means, 30 days"
        else:
            days, lev = o["dips"]
            n = int(((days >= M_dn(A.SEP[0])) & (days <= M_dn(A.SEP[1]))).sum())
            pumped = int(((g["dips_all"]["code"] == k) & g["dips_all"]["pump_remark"]
                          & (g["dips_all"]["reading_date"] >= "2026-09-01")
                          & (g["dips_all"]["reading_date"] <= "2026-09-30")).sum())
            basis = "%d static dips" % n + (" (%d pumped dip excluded)" % pumped if pumped else "")
        rows.append(dict(code=k, rn=b["rn"], name=b["name"], prop=b["prop"], owner=b["owner"], unit=b["unit"],
                         aquifer=b["aquifer"], cls=b["cls"], trigger=A.TRIGGER[b["cls"]], base_date=b["base_date"],
                         base_level=r2(o["baseline"]), level=r2(g["levels"][k]), change=r2(g["change"][k]),
                         drawdown=r1(g["drawdown"][k]), impaired="Yes" if g["impaired"][k] else "No",
                         offer=(OFFER_REF % offer_no[b["owner"]]) if g["impaired"][k] else "",
                         complaint=b["complaint"], basis=basis, fit=g["fits"][k]))
    return rows, offer_no


def M_dn(day):
    return (day - dt.date(1990, 1, 1)).days


# ------------------------------------------------------------------ the workbook

THIN = Side(style="thin", color="B7BEC4")
HEADFILL = PatternFill("solid", fgColor="DDE5DC")
YESFILL = PatternFill("solid", fgColor="F6E3D3")


def write_workbook(path, g, rows, offer_no):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Assessment"
    ws["A1"] = "Kelvale Mine: MG-02 bore assessments, September 2026"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("For the executive team, 26/10/2026. Assessment level: mean of the September 2026 static levels "
                "(MG-02 2.4). Prepared by %s, principal hydrogeologist, 16/10/2026." % AUTHOR)
    ws["A2"].font = Font(italic=True, size=9)
    cols = [("Bore", 10), ("Registered number", 11), ("Landholder's name", 17), ("Property", 15), ("Owner", 29),
            ("Aquifer", 15), ("Class", 15), ("Trigger (m)", 8), ("Baseline assessment", 11),
            ("Baseline level (m AHD)", 10), ("September 2026 level (m AHD)", 11), ("Change since baseline (m)", 10),
            ("Decline caused by the mine (m)", 11), ("Impaired capacity", 9), ("Offer", 9),
            ("Complaint received", 11), ("September level from", 30)]
    for j, (name, width) in enumerate(cols, 1):
        c = ws.cell(row=4, column=j, value=name)
        c.font = Font(bold=True, size=9)
        c.fill = HEADFILL
        c.alignment = Alignment(wrap_text=True, vertical="bottom", horizontal="center" if j >= 8 else "left")
        c.border = Border(bottom=THIN)
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.row_dimensions[4].height = 42
    r = 5
    for x in rows:
        vals = [x["code"], x["rn"], x["name"], x["prop"], x["owner"], UNIT_WORD[x["unit"]], x["cls"], x["trigger"],
                x["base_date"], x["base_level"], x["level"], x["change"], x["drawdown"], x["impaired"], x["offer"],
                x["complaint"], x["basis"]]
        for j, v in enumerate(vals, 1):
            if v in ("", None):
                continue
            c = ws.cell(row=r, column=j, value=v)
            c.font = Font(size=9, bold=(j == 14 and v == "Yes"))
            if j == 8:
                c.number_format = "0"
            elif j in (9, 16):
                c.number_format = "dd/mm/yyyy"
            elif j in (10, 11):
                c.number_format = "0.00"
            elif j == 12:
                c.number_format = "+0.00;-0.00;0.00"
            elif j == 13:
                c.number_format = "0.0"
            if j >= 8 and j not in (17,):
                c.alignment = Alignment(horizontal="right" if j < 14 else "center")
            if x["impaired"] == "Yes" and j in (13, 14):
                c.fill = YESFILL
        r += 1
    n_imp = sum(1 for x in rows if x["impaired"] == "Yes")
    ws.cell(row=r + 1, column=1, value="Bores with impaired capacity: %d, held by %d owners: %d make-good offers "
                                       "(MG-02 6.1, one agreement per owner)." % (n_imp, len(offer_no), len(offer_no))).font = Font(bold=True, size=9)
    ws.cell(row=r + 2, column=1, value="Change since baseline: September 2026 level less the baseline assessment "
                                       "level (negative is a fall). Decline caused by the mine: see Method.").font = Font(size=8, italic=True)
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = "A4:Q%d" % (r - 1)
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws2 = wb.create_sheet("Method")
    ws2["A1"] = "How the decline caused by the mine was estimated"
    ws2["A1"].font = Font(bold=True, size=12)
    text = [
        "Each bore's own response to daily rainfall less a fitted share of areal potential evaporation (a gamma-shaped "
        "response: gain, mean response time and shape fitted to the bore) was fitted on its readings from the first "
        "reading to 3 July 2022, before dewatering began on 4 July 2022. Shape, mean response time and the "
        "evaporation share were optimised continuously from a starting grid reaching 20 for the shape and 12 years "
        "for the mean; gain and level by least squares.",
        "Run forward on the rainfall and evaporation since, the fitted response gives the bore's September 2026 level "
        "from the climate alone; the decline caused by the mine is that level less the September 2026 assessment level.",
        "Readings: static dips reduced to the measuring point in force on the day (MG-02 2.1, 2.2), the 2012 to 2014 "
        "EIS rounds converted from feet, duplicate entries dropped; logged bores use daily means of hourly readings "
        "compensated with the site barometric logger, Queensland days (2.3).",
        "Checked against: the same fit on the three state reference bores (no drawdown) tracks them to September 2026; "
        "other response forms (delayed exponential, rainfall departure on each bore's own window, rainfall moving "
        "averages), calibration from May 2015, monthly steps, and a full-record fit with the take as a second stress "
        "give the same five bores.",
    ]
    for i, t in enumerate(text, 3):
        c = ws2.cell(row=i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws2.row_dimensions[i].height = max(62, 15 * math.ceil(len(t) / 95))
    hdr = ["Bore", "Readings fitted", "Mean response time (years)", "Shape", "Fit RMSE (m)"]
    for j, h in enumerate(hdr, 1):
        c = ws2.cell(row=9, column=j, value=h)
        c.font = Font(bold=True, size=9)
        c.fill = HEADFILL
        c.alignment = Alignment(wrap_text=True)
    for j, wdt in enumerate((12, 10, 13, 8, 11, 14, 14, 14), 1):
        ws2.column_dimensions[get_column_letter(j)].width = wdt
    for i, x in enumerate(rows, 10):
        fit = x["fit"]
        n, mean = fit["kernel"]
        vals = [x["code"], fit["n_obs"], round(mean / 365.25, 2), round(n, 1), round(fit["rmse"], 3)]
        fmts = [None, "0", "0.00", "0.0", "0.000"]
        for j, (v, f) in enumerate(zip(vals, fmts), 1):
            c = ws2.cell(row=i, column=j, value=v)
            c.font = Font(size=9)
            if f:
                c.number_format = f
    for i, t in enumerate(text, 3):
        ws2.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)
    ws2.page_setup.orientation = "portrait"
    ws2.page_setup.fitToWidth = 1
    ws2.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)
    PP.normalize_ooxml(path, AUTHOR, AUTHOR, "2026-10-14T05:12:00Z", "2026-10-16T02:47:00Z",
                       title="MG-02 bore assessments, September 2026")


# ------------------------------------------------------------------ the note

def _cell_shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _run(p, text, bold=False, size=None, italic=False, color=None):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def m(x, places=2):
    return ("%%.%df" % places) % abs(x)


def write_memo(path, g, rows, offer_no, ev):
    by = {x["code"]: x for x in rows}
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.5)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(9.5)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.04
    p = doc.add_paragraph()
    _run(p, "Tarrabool Resources  |  Kelvale Mine", size=8, color="5A6660")
    p = doc.add_paragraph()
    _run(p, "MEMO", bold=True, size=13)
    head = [("To", "Stephen Anderson, General Manager"), ("From", "Rebecca Quinn, Principal Hydrogeologist"),
            ("Copy", "Erika Martin, Environment Superintendent"), ("Date", "16 October 2026"),
            ("Subject", "MG-02 make good: four offers for five bores, for the executive team on 26 October")]
    for k, v in head:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        _run(p, "%s:\t" % k, bold=True)
        _run(p, v, bold=(k == "Subject"))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    imp = [x for x in rows if x["impaired"] == "Yes"]
    coal = [x for x in imp if x["unit"] == "coal"]
    alluv = [x for x in imp if x["unit"] == "alluvium"]
    owners = sorted(offer_no, key=lambda o: offer_no[o])
    p = doc.add_paragraph()
    _run(p, "I recommend we offer make-good agreements to %s owners for %s bores that the pit has drawn down past their "
            "MG-02 triggers at the September 2026 assessment level: " % (WORD[len(owners)], WORD[len(imp)]))
    parts = []
    for o in owners:
        bs = [x for x in imp if x["owner"] == o]
        parts.append("%s (%s, %s)" % (o, " and ".join("the %s %s" % (x["name"], x["code"]) for x in bs), bs[0]["prop"]))
    _run(p, "; ".join(parts) + ". Yarrowmere's two bores go in one agreement under MG-02 6.1, so that is %s offers." % WORD[len(owners)])

    t = doc.add_table(rows=1, cols=7)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    t.autofit = False
    hdr = ["Offer", "Owner", "Bore", "Aquifer", "Change since baseline (m)", "Decline caused by the mine (m)", "Trigger (m)"]
    widths = [Cm(1.7), Cm(4.9), Cm(3.4), Cm(2.4), Cm(2.0), Cm(2.1), Cm(1.3)]
    for i, wd in enumerate(widths):
        t.columns[i].width = wd
    for i, h in enumerate(hdr):
        c = t.rows[0].cells[i]
        c.text = ""
        _run(c.paragraphs[0], h, bold=True, size=8)
        _cell_shade(c, "DDE5DC")
        c.width = widths[i]
    for x in sorted(imp, key=lambda x: (x["offer"], x["code"])):
        cells = t.add_row().cells
        vals = [x["offer"], x["owner"], "%s, %s" % (x["name"], x["code"]), UNIT_WORD[x["unit"]], "%+.2f" % x["change"],
                "%.1f" % x["drawdown"], "%d" % x["trigger"]]
        for i, v in enumerate(vals):
            cells[i].text = ""
            para = cells[i].paragraphs[0]
            _run(para, v.replace("+", "") if i == 4 and v.startswith("+") else v.replace("-", "\u2212") if i == 4 else v, size=8)
            if i >= 4:
                para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            cells[i].width = widths[i]
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    _run(p, "Source: September 2026 assessment levels (MG-02 2.4) and baseline assessments in the landholder bore "
            "register; per-bore figures for all 18 bores in bore_impact_assessment.xlsx.", size=7.5, italic=True,
         color="4A544E")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    lo, hi = min(-x["change"] for x in coal), max(-x["change"] for x in coal)
    dlo, dhi = min(x["drawdown"] for x in coal), max(x["drawdown"] for x in coal)
    a1 = alluv[0]
    _run(p, "The coal-measure bores have fallen only %.1f to %.1f m since their baseline assessments, which is why they "
            "were never on our list. The coal measures answer rain over six years or more, and these bores were still "
            "rising on the 2019-20 and 2020-21 wet seasons when dewatering began on 4 July 2022 (about %.1f to %.1f m a "
            "year through 2021 and 2022). Each bore's own record, run forward on the rain since, puts it %.1f to %.1f m "
            "above where it stands now. The %s on %s is in the creek alluvium about a kilometre from the pit perimeter bores: "
            "%s m of its %s m fall is the mine's." % (lo, hi, EV["rise_lo"], EV["rise_hi"], dlo, dhi, a1["name"], a1["prop"],
                                                    "%.1f" % a1["drawdown"], m(a1["change"])))
    p = doc.add_paragraph()
    _run(p, "The four complaints do not meet the test. ", bold=True)
    _run(p, "Each bore did fall, but the dry seasons lowered it, not the pit. What we can tell each landholder:")
    lines = []
    for i, k in enumerate(g["complainants"]):
        x = by[k]
        if x["unit"] != "alluvium":
            lines.append((x, "your %s is %s m lower than at its May 2022 baseline; the ridge basalt drains quickly in dry "
                             "years and the four failed wet seasons account for all but %.1f m of the fall, well inside the "
                             "5 m trigger for the basalt." % (x["name"], m(x["change"]), x["drawdown"])))
        elif len([y for y, _ in lines if y["unit"] == "alluvium"]) == 0:
            lines.append((x, "your %s is %s m lower than at its June 2022 baseline; about %.1f m of that is the mine's and "
                             "the rest is the four failed wet seasons since 2022-23, much the fall the state's reference "
                             "bores 25 to 32 km from the mine show." % (x["name"], m(x["change"]), x["drawdown"])))
        elif len([y for y, _ in lines if y["unit"] == "alluvium"]) == 1:
            lines.append((x, "the %s has fallen %s m since June 2022, %.1f m of it from the mine; the creek alluvium has "
                             "been running down through four dry seasons on its own." % (x["name"], m(x["change"]), x["drawdown"])))
        else:
            lines.append((x, "the %s is down %s m on its June 2022 baseline and the mine's share is about %.1f m, short of "
                             "the 2 m trigger; the dry seasons did the rest." % (x["name"], m(x["change"]), x["drawdown"])))
    for x, text in lines:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        _run(p, "%s, %s (%s): " % (x["owner"], x["name"], x["prop"]), bold=True)
        _run(p, text)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    _run(p, "How the mine's share was separated. ", bold=True)
    _run(p, "For each of the 18 bores I fitted its own response to daily rainfall and evaporation on its record before "
            "4 July 2022 (the alluvium answers within months, the basalt within one to two years, the coal measures over "
            "six years or more), ran it forward on the rain and evaporation since, and took the mine's share as the gap "
            "between that level and the bore's September 2026 level. Fitted the same way, the three state reference "
            "bores, which the pit cannot reach, are tracked to within %.1f m through September 2026. The same five bores "
            "come out under every other response form and calibration period I tried, and with the take fitted as a "
            "second stress over the whole record. Measured from the baseline assessments alone, every alluvial bore "
            "but the two at Coolibah Springs and the ridge bore would pass the trigger and no coal-measure bore would." % EV["track"])
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    _run(p, "Levels are static readings (MG-02 2.1) reduced to the measuring point in force on the day (2.2); logged "
            "bores use daily means of readings compensated for air pressure on Queensland days (2.3). Changes are to the "
            "centimetre and the mine's share to 0.1 m; the trigger test is on unrounded figures.", size=7.5, italic=True, color="4A544E")
    doc.core_properties.author = AUTHOR
    doc.save(path)
    PP.normalize_ooxml(path, AUTHOR, AUTHOR, "2026-10-15T23:05:00Z", "2026-10-16T03:31:00Z",
                       title="MG-02 make good: four offers for five bores")


WORD = {1: "one", 2: "two", 3: "three", 4: "four", 5: "five", 6: "six"}
EV = {}


def main(target=None, out=None):
    target = target or os.path.join(TASK, "target")
    out = out or os.path.join(TASK, "golden")
    os.makedirs(out, exist_ok=True)
    g = G.solve(target, with_grid=False)
    g["dips_all"] = G.observations(target, pumped="keep")[4]
    ev = G.evidence(g)
    coal = [k for k in g["LH"] if g["obs"][k]["unit"] == "coal" and g["impaired"][k]]
    EV["rise_lo"] = math.floor(10 * min(ev["rise_2021_22"][k] for k in coal)) / 10
    EV["rise_hi"] = math.ceil(10 * max(ev["rise_2021_22"][k] for k in coal)) / 10
    EV["track"] = math.ceil(10 * max(ev["ref_tracking_max_abs"].values())) / 10
    rows, offer_no = figures(g)
    write_workbook(os.path.join(out, WORKBOOK), g, rows, offer_no)
    write_memo(os.path.join(out, MEMO), g, rows, offer_no, ev)
    return g, rows


if __name__ == "__main__":
    main()
