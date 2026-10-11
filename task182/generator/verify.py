"""Independent verifier. Reads only target/ and golden/, imports nothing from the generator, and recomputes on its
own code path: every bore's September 2026 assessment level and change since its baseline, the naive rungs, and
the main call (a monthly gamma response to rainfall and evaporation fitted before 4 July 2022, written afresh), the
decline column (a daily gamma response searched continuously, within the stated bands), then matches the golden
workbook and note."""
import csv
import datetime as dt
import math
import os
import re
import sys
from collections import defaultdict

import numpy as np
import openpyxl
import pyarrow.parquet as pq

TASK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
D = dt.date
EPOCH = D(1990, 1, 1)


def ddate(v):
    if v is None:
        return None
    return v.date() if isinstance(v, dt.datetime) else v


def register(target):
    wb = openpyxl.load_workbook(os.path.join(target, "landholder_bore_register.xlsx"), read_only=True, data_only=True)

    def table(name):
        rows = list(wb[name].iter_rows(min_row=4, values_only=True))
        return [dict(zip(rows[0], r)) for r in rows[1:] if r[0] and str(r[0]).startswith("KV-")]
    lh = {r["Bore"]: r for r in table("Landholder bores")}
    net = {r["Bore"]: r for r in table("Network bores")}
    mp = defaultdict(list)
    for r in table("Measuring points"):
        mp[r["Bore"]].append((ddate(r["In force from"]), ddate(r["In force to"]), r["MP elevation (m AHD)"]))
    return lh, net, mp


def mp_at(mp, code, day):
    hit = [e for f, t, e in mp[code] if f <= day and (t is None or day <= t)]
    assert len(hit) == 1, (code, day)
    return hit[0]


def kind(aq):
    s = aq.lower()
    return "alluvium" if "alluv" in s else "basalt" if "basalt" in s else "coal"


def main(target, golden):
    lh, net, mp = register(target)
    names = {}
    for code, r in list(lh.items()) + list(net.items()):
        rn = str(r["Registered number"])
        for s in (code, code[3:], rn, "RN" + rn):
            names[re.sub(r"[^A-Z0-9]", "", s.upper())] = code
    # manual dips: static, distinct entries, metres, reduced level with the measuring point in force on the day
    seen, dips = set(), defaultdict(list)
    with open(os.path.join(target, "water_level_dips_2012-2026.csv"), newline="") as fh:
        for r in csv.DictReader(fh):
            key = tuple(r[c] for c in ("site", "reading_date", "reading_time", "dtw", "units", "remarks", "officer"))
            code = names[re.sub(r"[^A-Z0-9]", "", r["site"].upper())]
            key = (code,) + key[1:]
            if key in seen:
                continue
            seen.add(key)
            rem = r["remarks"].lower()
            if "pump" in rem and "pump off" not in rem:
                continue
            day = D.fromisoformat(r["reading_date"])
            depth = float(r["dtw"]) * (0.3048 if r["units"] == "ft" else 1.0)
            dips[code].append((day, mp_at(mp, code, day) - depth))
    # loggers: hourly pressure less the site barometric pressure for the same hour, on Queensland days
    baro = {}
    with open(os.path.join(target, "barometric_logger_2015-2026.csv"), newline="") as fh:
        for r in csv.DictReader(fh):
            t = dt.datetime.strptime(r["timestamp_utc"], "%Y-%m-%dT%H:%M:%SZ")
            baro[int((t - dt.datetime(1970, 1, 1)).total_seconds())] = float(r["pressure_kpa"])
    bk = np.array(sorted(baro))
    bv = np.array([baro[k] for k in bk])
    tab = pq.read_table(os.path.join(target, "logger_pressure_2015-2026.parquet"))
    ser = np.array(tab.column("logger_serial").to_pylist())
    secs = tab.column("timestamp_utc").to_numpy().astype("datetime64[s]").astype(np.int64)
    kpa = tab.column("pressure_kpa").to_numpy()
    logged = defaultdict(dict)
    with open(os.path.join(target, "logger_installations.csv"), newline="") as fh:
        for r in csv.DictReader(fh):
            if not r["sensor_depth_below_mp_m"]:
                continue
            sel = ser == r["logger_serial"]
            s, p = secs[sel], kpa[sel]
            j = np.searchsorted(bk, s)
            good = (j < len(bk)) & (bk[np.minimum(j, len(bk) - 1)] == s)
            s, p, b = s[good], p[good], bv[j[good]]
            local_day = (s + 36000) // 86400                      # days since 1970 on Queensland time
            sums = defaultdict(lambda: [0.0, 0])
            for dday, col in zip(local_day.tolist(), ((p - b) / 9.80665).tolist()):
                sums[dday][0] += col
                sums[dday][1] += 1
            code = r["site"]
            depth = float(r["sensor_depth_below_mp_m"])
            for dday, (tot, n) in sums.items():
                day = D(1970, 1, 1) + dt.timedelta(days=dday)
                logged[code][day] = mp_at(mp, code, day) - depth + tot / n
    sep_lo, sep_hi = D(2026, 9, 1), D(2026, 9, 30)

    def september(code):
        if logged.get(code):
            v = [x for d_, x in logged[code].items() if sep_lo <= d_ <= sep_hi]
            assert len(v) == 30, code
        else:
            v = [x for d_, x in dips[code] if sep_lo <= d_ <= sep_hi]
        return sum(v) / len(v)

    def readings(code):
        """Logger daily means from the logger's first day, static dips before it."""
        out = []
        first = min(logged[code]) if logged.get(code) else None
        for d_, x in dips[code]:
            if first is None or d_ < first:
                out.append((d_, x))
        if first is not None:
            out += sorted(logged[code].items())
        return sorted(out)

    def mean_between(code, a, b):
        v = [x for d_, x in readings(code) if a <= d_ <= b]
        return sum(v) / len(v)

    level = {c: september(c) for c in list(lh) + list(net)}
    base = {c: mp_at(mp, c, ddate(r["Baseline assessment"])) - float(r["Baseline static water level (m below MP)"])
            for c, r in lh.items()}
    change = {c: level[c] - base[c] for c in lh}
    trig = {c: 2.0 if r["Aquifer class"] == "unconsolidated" else 5.0 for c, r in lh.items()}
    unit = {c: kind(r["Aquifer"]) for c, r in list(lh.items()) + list(net.items())}

    def designated(q):
        return sorted(c for c in lh if q[c] > trig[c])

    # rung 1 and 1b: the decline from a reference level
    r1 = {c: base[c] - level[c] for c in lh}
    r1j = {c: mean_between(c, D(2022, 6, 1), D(2022, 6, 30)) - level[c] for c in lh}
    pre_month = {}
    for c in list(lh) + list(net):
        mm = defaultdict(list)
        for d_, x in readings(c):
            if d_ <= D(2022, 6, 30):
                mm[(d_.year, d_.month)].append(x)
        pre_month[c] = {k: sum(v) / len(v) for k, v in mm.items()}
    r1m = {c: sum(pre_month[c].values()) / len(pre_month[c]) - level[c] for c in lh}
    refs = [c for c, r in net.items() if "reference" in r["Type"]]
    rch = {c: mean_between(c, D(2022, 4, 1), D(2022, 6, 30)) - level[c] for c in refs}
    pooled = sum(rch.values()) / len(rch)
    by_unit = {u: [rch[c] for c in refs if unit[c] == u] for u in ("alluvium", "basalt")}
    matched = {u: sum(v) / len(v) for u, v in by_unit.items()}
    r2p = {c: r1[c] - pooled for c in lh}
    r2m = {c: r1[c] - matched["alluvium" if unit[c] == "alluvium" else "basalt"] for c in lh}
    # the climate, monthly
    rain = defaultdict(float)
    pet = defaultdict(float)
    with open(os.path.join(target, "daily_climate_1990-2026.csv"), newline="") as fh:
        for r in csv.DictReader(fh):
            k = (int(r["date"][:4]) - 1990) * 12 + int(r["date"][5:7]) - 1
            rain[k] += float(r["rain_mm"])
            pet[k] += float(r["morton_apet_mm"])
    nm = max(rain) + 1
    P = np.array([rain[i] for i in range(nm)])
    E = np.array([pet[i] for i in range(nm)])
    P -= P.mean()
    E -= E.mean()

    def mi(d_):
        return (d_.year - 1990) * 12 + d_.month - 1

    cal = mi(D(2022, 6, 1))
    sep = mi(D(2026, 9, 1))
    # rung 2b: a straight line through the monthly levels, extrapolated to mid-September 2026
    r2b = {}
    for start in (2017, 2012):
        q = {}
        for c in lh:
            ks = sorted(k for k in pre_month[c] if k[0] >= start)
            x = np.array([mi(D(y, m, 15)) for y, m in ks], float)
            y = np.array([pre_month[c][k] for k in ks])
            slope, icpt = np.polyfit(x, y, 1)
            q[c] = icpt + slope * (sep + 14.0 / 30.4375) - level[c]
        r2b[start] = q
    # rung 3: one window of summed rainfall for every bore (a slope below zero read as no response)
    csum = np.concatenate([[0.0], np.cumsum(P)])
    r3 = {}
    for w in (12, 36):
        win = np.array([csum[t + 1] - csum[t + 1 - w] if t + 1 - w >= 0 else np.nan for t in range(nm)])
        q = {}
        for c in lh:
            ks = sorted(pre_month[c])
            x = np.array([win[mi(D(y, m, 1))] for y, m in ks])
            y = np.array([pre_month[c][k] for k in ks])
            ok_ = ~np.isnan(x)
            slope, icpt = np.polyfit(x[ok_], y[ok_], 1)
            q[c] = (icpt + slope * win[sep] if slope > 0 else y[ok_].mean()) - level[c]
        r3[w] = q
    # the main call: each bore's monthly gamma response to rainfall less f x evaporation, fitted to June 2022
    kern = []
    for n in (1.0, 1.5, 2.0, 2.5, 3.0, 4.0, 5.0, 7.0, 10.0, 14.0, 18.0, 24.0):
        for mean_m in np.exp(np.linspace(math.log(0.6), math.log(150.0), 45)):
            a = mean_m / n
            L = int(min(nm, mean_m * 6 + 48))
            t = np.arange(L) + 0.5
            w = np.exp((n - 1) * np.log(t) - t / a - math.lgamma(n) - n * math.log(a))
            w /= w.sum()
            kern.append(((n, mean_m), np.convolve(P, w)[:nm], np.convolve(E, w)[:nm]))
    fs = np.linspace(0.0, 1.0, 41)
    dd = {}
    for c in lh:
        ks = sorted(pre_month[c])
        idx = np.array([mi(D(y, m, 1)) for y, m in ks])
        y = np.array([pre_month[c][k] for k in ks])
        best = None
        yc = y - y.mean()
        for meta, xp, xe in kern:
            z = xp[idx][None, :] - fs[:, None] * xe[idx][None, :]
            zc = z - z.mean(1, keepdims=True)
            vz = (zc * zc).sum(1)
            gain = (zc @ yc) / vz
            sse = (yc * yc).sum() - gain * (zc @ yc)
            sse = np.where(gain > 0, sse, np.inf)
            j = int(np.argmin(sse))
            if best is None or sse[j] < best[0]:
                best = (sse[j], gain[j], y.mean() - gain[j] * z[j].mean(), xp, xe, fs[j])
        sse, gain, c0, xp, xe, f = best
        dd[c] = c0 + gain * (xp[sep] - f * xe[sep]) - level[c]
    call = designated(dd)
    # the decline column: each bore's daily gamma response, shape, mean and f searched continuously (a compass search
    # in log shape and log mean from the best of a wide start grid; f and the gain by constrained least squares),
    # fitted on its static dips and logger daily means to 3 July 2022
    rain_d, pet_d = [], []
    with open(os.path.join(target, "daily_climate_1990-2026.csv"), newline="") as fh:
        for r in csv.DictReader(fh):
            rain_d.append(float(r["rain_mm"]))
            pet_d.append(float(r["morton_apet_mm"]))
    Pd = np.array(rain_d) - np.mean(rain_d)
    Ed = np.array(pet_d) - np.mean(pet_d)
    nd = len(Pd)
    nfft = 1 << (2 * nd).bit_length()
    FP, FE = np.fft.rfft(Pd, nfft), np.fft.rfft(Ed, nfft)
    tt = np.arange(nd) + 0.5

    def responses(n, mean_d):
        a = mean_d / n
        w = np.exp((n - 1) * np.log(tt) - tt / a - math.lgamma(n) - n * math.log(a))
        w /= w.sum()
        fw = np.fft.rfft(w, nfft)
        return np.fft.irfft(FP * fw, nfft)[:nd], np.fft.irfft(FE * fw, nfft)[:nd]

    def lsq(xp, xe, y):
        best = (np.inf, None)
        X = np.column_stack([np.ones_like(y), xp, xe])
        sol = np.linalg.lstsq(X, y, rcond=None)[0]
        if sol[1] > 0 and -sol[1] <= sol[2] <= 0:
            res = y - X @ sol
            return float(res @ res), (sol[0], sol[1], -sol[2] / sol[1])
        for f in (0.0, 1.0):
            z = xp - f * xe
            A_ = np.column_stack([np.ones_like(y), z])
            c0, gn = np.linalg.lstsq(A_, y, rcond=None)[0]
            if gn > 0:
                res = y - A_ @ np.array([c0, gn])
                if res @ res < best[0]:
                    best = (float(res @ res), (c0, gn, f))
        return best

    cal_end = D(2022, 7, 3)
    sep_days = np.array([(D(2026, 9, 1) - EPOCH).days + i for i in range(30)])
    dd_daily = {}
    for c in lh:
        pts = [((d_ - EPOCH).days, x) for d_, x in readings(c) if d_ <= cal_end]
        ix = np.array([p_[0] for p_ in pts])
        y = np.array([p_[1] for p_ in pts])
        cand = []
        for n in (1.0, 1.5, 2.0, 3.0, 5.0, 8.0, 12.0, 16.0, 20.0):
            for mean_d in np.exp(np.linspace(math.log(30.0), math.log(12 * 365.25), 32)):
                xp, xe = responses(n, mean_d)
                cand.append((lsq(xp[ix], xe[ix], y)[0], math.log(n), math.log(mean_d)))
        cand.sort()
        best = None
        for sse0, u, v in cand[:3]:
            step_ = 0.2
            cur = sse0
            while step_ > 1e-4:
                moved = False
                for du, dv in ((step_, 0), (-step_, 0), (0, step_), (0, -step_)):
                    nu, nv = u + du, v + dv
                    if math.exp(nu) < 1.0:
                        continue
                    xp, xe = responses(math.exp(nu), math.exp(nv))
                    sc = lsq(xp[ix], xe[ix], y)[0]
                    if sc < cur:
                        cur, u, v, moved = sc, nu, nv, True
                        break
                if not moved:
                    step_ /= 2
            if best is None or cur < best[0]:
                best = (cur, u, v)
        _, u, v = best
        xp, xe = responses(math.exp(u), math.exp(v))
        _, (c0, gn, f) = lsq(xp[ix], xe[ix], y)
        dd_daily[c] = float(np.mean(c0 + gn * (xp[sep_days] - f * xe[sep_days]))) - level[c]
    # against the golden workbook and note
    wb = openpyxl.load_workbook(os.path.join(golden, "bore_impact_assessment.xlsx"), data_only=True)
    ws = wb["Assessment"]
    head = [x.value for x in ws[4]]
    rows = {r[0]: dict(zip(head, r)) for r in ws.iter_rows(min_row=5, max_row=22, values_only=True)}
    assert sorted(rows) == sorted(lh), "workbook bores"
    for c in lh:
        r = rows[c]
        assert r["September 2026 level (m AHD)"] == round(level[c], 2), (c, r["September 2026 level (m AHD)"], level[c])
        assert r["Change since baseline (m)"] == round(change[c], 2), (c, r["Change since baseline (m)"], change[c])
        assert (r["Impaired capacity"] == "Yes") == (c in call), c
        assert abs(r["Decline caused by the mine (m)"] - dd[c]) <= 0.6, (c, r["Decline caused by the mine (m)"], dd[c])
        band = 0.4 if c == "KV-LB06" else 0.3
        assert abs(r["Decline caused by the mine (m)"] - dd_daily[c]) <= band, (c, r["Decline caused by the mine (m)"],
                                                                                dd_daily[c])
    owners = defaultdict(list)
    for c in call:
        owners[lh[c]["Owner"]].append(c)
    from docx import Document
    doc = Document(os.path.join(golden, "make_good_recommendation.docx"))
    memo = "\n".join(p.text for p in doc.paragraphs) + "\n".join(x.text for t in doc.tables for r in t.rows for x in r.cells)
    words = {3: "three", 4: "four", 5: "five"}
    assert "%s offers" % words[len(owners)] in memo, "offer count in the note"
    for o, cs in owners.items():
        assert o in memo and all(x in memo for x in cs), o
    complainants = [c for c, r in lh.items() if r["Complaint received (MG-02 7.1)"]]
    for c in complainants:
        assert c not in call and lh[c]["Landholder's name for the bore"] in memo, c
    rungs = {"R1": designated(r1), "R1b June": designated(r1j), "R1b mean": designated(r1m), "R2 pooled": designated(r2p),
             "R2 matched": designated(r2m), "R2b 2017": designated(r2b[2017]), "R2b 2012": designated(r2b[2012]),
             "R3 12": designated(r3[12]), "R3 36": designated(r3[36])}
    for name, s in rungs.items():
        assert s != call, name
    print("verify: OK. %d bores impaired (%s), %d offers; 18 levels and 18 changes match the workbook exactly; "
          "declines within their bands on its own daily fit (largest gap %.2f m; monthly coarse fit %.2f m); "
          "naive rungs: %s" % (
              len(call), ", ".join(call), len(owners),
              max(abs(rows[c]["Decline caused by the mine (m)"] - dd_daily[c]) for c in lh),
              max(abs(rows[c]["Decline caused by the mine (m)"] - dd[c]) for c in lh),
              "; ".join("%s %s" % (k, ",".join(v)) for k, v in rungs.items())))
    return dict(call=call, level=level, change=change, dd=dd, dd_daily=dd_daily, rungs=rungs)


if __name__ == "__main__":
    main(os.path.join(TASK, "target"), os.path.join(TASK, "golden"))
