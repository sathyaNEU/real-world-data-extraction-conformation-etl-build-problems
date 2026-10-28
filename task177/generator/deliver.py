"""The two golden deliverables: the pricing workbook and the purchase note with its one chart."""
from __future__ import annotations

import datetime as dt
import io
import os
import re
import zipfile

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.shared import Pt, Cm, RGBColor  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
import openpyxl  # noqa: E402
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side  # noqa: E402

import docs as DOCS  # noqa: E402
import figures as FG  # noqa: E402
import params as P  # noqa: E402

XLSX = "pool_90dpd_forecast.xlsx"
DOCX = "pool_purchase_note.docx"
WHEN = (2027, 1, 15, 17, 40, 0)
DESK = b"Tamralipta Bank Ltd"
INK = "#1f2933"
BAR = "#2f6f9f"


def fmt_int(x):
    return f"{int(x):,}"


def workbook(F, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Forecast"
    bold = Font(bold=True)
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="2F6F9F")
    thin = Side(style="thin", color="BBBBBB")
    ws["A1"] = "Pool PVM/2026-12 (Prachi Valley Microfinance): loans 90 or more days past due on 30 June 2027"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Cut-off 31 December 2026; forecast from the assignor's collection records, July 2023 to December 2026"
    ws["A4"], ws["B4"] = "Forecast, loans 90+ days past due on 30 June 2027 (nearest 100)", F["call"]
    ws["A5"], ws["B5"] = "Before rounding", F["call_unrounded"]
    ws["A6"], ws["B6"] = "Loans in the pool", F["pool_loans"]
    ws["A7"], ws["B7"] = "Share of the pool's loans (per cent)", F["call_share_pct"]
    for c in ("A4", "B4"):
        ws[c].font = bold
    hdr = ["Members the group paid for at its last meeting on or before the cut-off", "Loans in the pool",
           "Loans reaching 90+ days within six months, July 2023 to June 2026 (per cent)",
           "Expected loans 90+ on 30 June 2027"]
    ws.append([])
    ws.append(hdr)
    r0 = ws.max_row
    for j in range(1, 5):
        ws.cell(r0, j).font = head
        ws.cell(r0, j).fill = fill
        ws.cell(r0, j).alignment = Alignment(wrap_text=True, vertical="top")
    for k in range(4):
        ws.append([FG.STATE_NAMES[k].capitalize(), F["state_loans_%d" % k], F["state_rate_pct_%d" % k],
                   F["state_expected_%d" % k]])
    ws.append(["Total", F["pool_loans"], None, F["call_unrounded"]])
    for row in ws.iter_rows(min_row=r0 + 1, max_row=ws.max_row):
        for c in row:
            c.border = Border(bottom=thin)
        row[1].number_format = "#,##0"
        row[2].number_format = "0.00"
        row[3].number_format = "0.0"
    ws["B5"].number_format = "0.0"
    ws["B6"].number_format = "#,##0"
    ws["B7"].number_format = "0.0"
    ws.freeze_panes = "A%d" % (r0 + 1)
    ws.cell(ws.max_row, 1).font = bold
    ws.append([])
    ws.append(["The group's count is the number of its members' instalments paid at the meeting by another member "
               "(Kshetra GRP_SUPP, CollectPro GS). Rates are read at each month end on the loans current then and "
               "disbursed at least 91 days earlier, by the group's count at its last meeting on or before the month "
               "end, and the outcome 181 days later under the assignor's days-past-due rule."])
    ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=4)
    ws.row_dimensions[ws.max_row].height = 58
    ws.column_dimensions["A"].width = 46
    for col in "BCD":
        ws.column_dimensions[col].width = 22
    ws.row_dimensions[r0].height = 62
    h = wb.create_sheet("History")
    h.append(["Members the group paid for at its last meeting", "Loan-months at the month ends",
              "Reached 90+ days within six months", "Rate (per cent)"])
    for j in range(1, 5):
        h.cell(1, j).font = head
        h.cell(1, j).fill = fill
        h.cell(1, j).alignment = Alignment(wrap_text=True, vertical="top")
    for k in range(4):
        n = F["state_hist_loanmonths_%d" % k]
        h.append([FG.STATE_NAMES[k].capitalize(), n, F["state_hist_events_%d" % k], F["state_rate_pct_%d" % k]])
    for row in h.iter_rows(min_row=2, max_row=h.max_row):
        row[1].number_format = "#,##0"
        row[2].number_format = "#,##0"
        row[3].number_format = "0.00"
    h.freeze_panes = "A2"
    h.append([])
    h.append(["Month ends used", "%s to %s" % (F["hist_first_asof"], F["hist_last_asof"])])
    h.column_dimensions["A"].width = 40
    for col in "BCD":
        h.column_dimensions[col].width = 20
    h.row_dimensions[1].height = 46
    g = wb.create_sheet("Guardrails")
    g["A1"] = "Pool borrowers in breach of CPN-MF-07 on the bureau report at 31 December 2026"
    g["A1"].font = bold
    g.append([])
    g.append(["Branch code", "Branch", "Borrowers in breach", "Pool principal outstanding on their loans (INR lakh)"])
    for j in range(1, 5):
        g.cell(3, j).font = head
        g.cell(3, j).fill = fill
        g.cell(3, j).alignment = Alignment(wrap_text=True, vertical="top")
    for b in FG.BRANCH_ORDER:
        g.append([b, FG.BRANCH_NAME[b], F["ask_breach_%s" % b], F["ask_lakh_%s" % b]])
    g.append(["Total", None, F["ask_breach_total"], F["ask_lakh_total"]])
    g.cell(g.max_row, 1).font = bold
    for row in g.iter_rows(min_row=4, max_row=g.max_row):
        row[2].number_format = "#,##0"
        row[3].number_format = "0.0"
    g.freeze_panes = "A4"
    g.append([])
    g.append(["Active JLG and MFI-IND accounts; a lending institution counted once (Salandi Microfinance MFI0447 with "
              "Salandi Small Finance Bank SFB0029; Gandhamardan Gramya Bank with its BC portfolio BNK0115); a loan counted "
              "once at its latest report, a Nirmala Small Finance Bank loan reported under its pre-migration and its "
              "new account number taken at the new number's report."])
    g.cell(g.max_row, 1).alignment = Alignment(wrap_text=True)
    g.merge_cells(start_row=g.max_row, start_column=1, end_row=g.max_row, end_column=4)
    g.row_dimensions[g.max_row].height = 58
    g.row_dimensions[3].height = 44
    g.column_dimensions["A"].width = 14
    g.column_dimensions["B"].width = 18
    g.column_dimensions["C"].width = 16
    g.column_dimensions["D"].width = 30
    nt = wb.create_sheet("Sources")
    for line in [
        ["Sources and conventions"],
        ["As of", "Cut-off 31 December 2026; data room export of 8 January 2027"],
        ["Pool", "The loans in the assignor's pool tape at the cut-off (term sheet clause 3)"],
        ["Days past due", "Assignor's JLG operating policy v6.1, section 6: from the due date of the oldest instalment "
                          "not paid in full"],
        ["History", "Collections July 2023 to December 2026, CollectPro and Kshetra joined through the migration "
                    "map; month ends with 181 days of follow-up in the record"],
        ["Basis", "Members' repayment from 1 January 2027 taken as in the assignor's experience before the cyclone "
                  "of 16 October 2026 (term sheet clause 12)"],
        ["Guardrails", "CPN-MF-07 sections 2 and 3, on the bureau's portfolio review at the cut-off"],
        ["Rounding", "Forecast to the nearest 100 loans; principal in INR lakh to one decimal"],
    ]:
        nt.append(line)
    nt["A1"].font = bold
    nt.column_dimensions["A"].width = 18
    nt.column_dimensions["B"].width = 110
    wb.properties.creator = "Direct Assignment desk"
    wb.properties.lastModifiedBy = "Direct Assignment desk"
    wb.properties.title = "Pool PVM/2026-12 forecast"
    wb.save(path)
    DOCS.fixed_xlsx(path, when=WHEN, company=DESK)


def chart(F):
    labels = []
    vals = []
    for k in range(4):
        labels.append("%s\n(%s loans)" % (FG.STATE_NAMES[k].capitalize(), fmt_int(F["state_loans_%d" % k])))
        vals.append(F["state_expected_%d" % k])
    fig, ax = plt.subplots(figsize=(6.4, 3.1), dpi=200)
    y = list(range(len(vals)))[::-1]
    ax.barh(y, vals, height=0.52, color=BAR)
    for yi, v in zip(y, vals):
        ax.text(v + max(vals) * 0.015, yi, "%.1f" % v, va="center", ha="left", fontsize=8.5, color=INK)
    ax.set_yticks(y)
    ax.set_yticklabels(labels, fontsize=8.5, color=INK)
    ax.set_xlabel("Expected loans 90+ days past due on 30 June 2027", fontsize=8.5, color=INK)
    ax.set_title("Most of the expected 90-day loans sit in groups paying for one or two members", fontsize=9.5,
                 color=INK, loc="left")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    for s in ("left", "bottom"):
        ax.spines[s].set_color("#9aa5b1")
    ax.tick_params(axis="x", labelsize=8, colors="#52606d")
    ax.tick_params(axis="y", length=0)
    ax.xaxis.grid(True, color="#e4e7eb", linewidth=0.6)
    ax.set_axisbelow(True)
    ax.set_xlim(0, max(vals) * 1.15)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None, "Creation Time": None})
    plt.close(fig)
    return buf.getvalue()


def _field(paragraph, instr):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    r = paragraph.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText")
    i.set(qn("xml:space"), "preserve")
    i.text = instr
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r._r.append(b)
    r._r.append(i)
    r._r.append(e)


def note(F, path):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    fp = sec.footer.paragraphs[0]
    fp.text = "Tamralipta Bank Ltd, Direct Assignment desk. Internal. Pool PVM/2026-12. Page "
    _field(fp, "PAGE")
    fp.runs[0].font.size = Pt(8)
    memo = d.add_table(rows=5, cols=2)
    for i, (a, b) in enumerate([("To", "Credit Committee"), ("From", "Direct Assignment desk"),
                                ("Date", "15 January 2027"), ("Ref", "DA/PVM/2026-12"),
                                ("Subject", "Pool PVM/2026-12 from Prachi Valley Microfinance: six-month delinquency "
                                            "forecast and guardrail screen")]):
        memo.cell(i, 0).text = a
        memo.cell(i, 1).text = b
        memo.cell(i, 0).paragraphs[0].runs[0].bold = True
        memo.cell(i, 0).width = Cm(2.6)
        memo.cell(i, 1).width = Cm(14)
    d.add_paragraph()
    t = d.add_paragraph()
    r = t.add_run("About %s of the pool's loans will be 90 or more days past due on 30 June 2027" % fmt_int(F["call"]))
    r.bold = True
    r.font.size = Pt(13)
    p = d.add_paragraph()
    p.add_run("We expect about %s of the pool's %s loans to be 90 or more days past due on 30 June 2027 (%.1f "
              "before rounding, %.1f per cent of the loans)." % (fmt_int(F["call"]), fmt_int(F["pool_loans"]),
                                                                F["call_unrounded"], F["call_share_pct"]))
    p.add_run("1").font.superscript = True
    p.add_run(" Every loan in the pool is current at the cut-off, but some are current only because another member "
              "of the group paid the instalment at the centre meeting. At its last meeting on or before the cut-off, %s of "
              "the pool's loans sat in groups paying for one member and %s in groups paying for two or more. In the "
              "assignor's record a group carrying one member mostly keeps going, and %.2f per cent of its loans reach "
              "90 days within six months; a group carrying two at once often stops, and then every member's loan "
              "falls behind, the members who were paying included: %.2f per cent. The forecast applies the record's "
              "rate for each count, read at the month ends from July 2023 to June 2026, to the pool's loans at that "
              "count." % (fmt_int(F["state_loans_1"]), fmt_int(F["state_loans_2plus"]), F["state_rate_pct_1"],
                          F["state_rate_pct_2"]))
    tab = d.add_table(rows=1, cols=4)
    tab.style = "Light List"
    hdr = ["Members the group paid for at its last meeting", "Pool loans", "Six-month rate in the record (%)",
           "Expected 90+ on 30 June 2027"]
    for j, h in enumerate(hdr):
        tab.rows[0].cells[j].text = h
    for k in range(4):
        cells = tab.add_row().cells
        cells[0].text = FG.STATE_NAMES[k].capitalize()
        cells[1].text = fmt_int(F["state_loans_%d" % k])
        cells[2].text = "%.2f" % F["state_rate_pct_%d" % k]
        cells[3].text = "%.1f" % F["state_expected_%d" % k]
    cells = tab.add_row().cells
    cells[0].text = "Total"
    cells[1].text = fmt_int(F["pool_loans"])
    cells[3].text = "%.1f" % F["call_unrounded"]
    src = d.add_paragraph("Source: assignor's collections, CollectPro (July 2023 to March 2025) and Kshetra (April "
                          "2025 to December 2026); pool tape at 31 December 2026. The rate for three or more rests on "
                          "%d loan-months." % F["state_hist_loanmonths_3"])
    src.runs[0].font.size = Pt(8.5)
    buf = io.BytesIO(chart(F))
    d.add_picture(buf, width=Cm(15.5))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = d.add_paragraph()
    p.add_run("Guardrails. ").bold = True
    p.add_run("On the bureau report at the cut-off, %s of the pool's borrowers were in breach of CPN-MF-07, with "
              "INR %.1f lakh of pool principal outstanding on their loans; the workbook gives the branch split. "
              "Nirmala Small Finance Bank's pre-migration account numbers still show active on the bureau beside the "
              "new ones, and the bureau's own lender count treats Salandi's two member codes and Gandhamardan's BC "
              "portfolio as separate lenders, so the bureau's summary overstates breaches."
              % (fmt_int(F["ask_breach_total"]), F["ask_lakh_total"]))
    p = d.add_paragraph()
    p.add_run("For the decision. ").bold = True
    p.add_run("Price and size the pool on about %s loans 90 or more days past due on 30 June 2027, most of them in "
              "the Kendrapara and Jagatsinghpur districts, not on the assignor's portfolio report, which counts a "
              "co-member's payment as collected." % fmt_int(F["call"]))
    fn = d.add_paragraph()
    fr = fn.add_run("1")
    fr.font.superscript = True
    fr.font.size = Pt(8)
    fn.add_run(" Days past due as the assignor's JLG operating policy defines them (version 6.1, section 6): from the "
               "due date of the oldest instalment not paid in full.").font.size = Pt(8)
    cp = d.core_properties
    cp.author = "Direct Assignment desk"
    cp.last_modified_by = "Direct Assignment desk"
    cp.title = "Pool PVM/2026-12 purchase note"
    cp.comments = ""
    cp.created = dt.datetime(*WHEN)
    cp.modified = dt.datetime(*WHEN)
    cp.revision = 1
    d.save(path)
    _fix_docx(path)


def _fix_docx(path):
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            if info.filename == "docProps/app.xml":
                data = (b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Properties xmlns="http://schemas.'
                        b'openxmlformats.org/officeDocument/2006/extended-properties" xmlns:vt="http://schemas.'
                        b'openxmlformats.org/officeDocument/2006/docPropsVTypes"><Template>Normal.dotm</Template>'
                        b'<TotalTime>34</TotalTime><Application>Microsoft Office Word</Application><DocSecurity>0'
                        b'</DocSecurity><Company>Tamralipta Bank Ltd</Company><AppVersion>16.0000</AppVersion>'
                        b'</Properties>')
            zi = zipfile.ZipInfo(info.filename, date_time=WHEN)
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, data)


def build(F, folder):
    os.makedirs(folder, exist_ok=True)
    workbook(F, os.path.join(folder, XLSX))
    note(F, os.path.join(folder, DOCX))
