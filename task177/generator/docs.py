"""task177 documents: the deduction agreements register and the data dictionary (Excel), the Libranza Empresarial
product rules and the cooperative's 2025 annual return to its supervisor (PDF). Their text is held here as data so
checks.py can quote it and verify.py can read it back."""
from __future__ import annotations

import datetime as dt
import os

import numpy as np
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                        # noqa: E402
from reportlab.lib.pagesizes import A4                                                  # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                         # noqa: E402
from reportlab.lib.units import mm                                                      # noqa: E402
from reportlab.pdfbase import pdfmetrics                                                # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                            # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import common as C          # noqa: E402
import postprocess as PP    # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/"
for _n, _f in (("Carlito", "crosextra/Carlito-Regular.ttf"), ("Carlito-Bold", "crosextra/Carlito-Bold.ttf"),
               ("LibSans", "liberation/LiberationSans-Regular.ttf"), ("LibSans-Bold", "liberation/LiberationSans-Bold.ttf")):
    pdfmetrics.registerFont(TTFont(_n, FONT_DIR + _f))
INK = colors.HexColor("#1b2430")
GREY = colors.HexColor("#5d6875")
ACCENT = colors.HexColor("#1f5b4a")
RULE = colors.HexColor("#c9d1d6")

# --------------------------------------------------------------------------------------------- the product rules
RULES_DATE = dt.date(2026, 11, 20)
RULES_TITLE = "Libranza Empresarial: product rules"
RULES_SUB = ("Coopservir, Credit Directorate. Version 2.0, approved by the Board of Directors on 20 November 2026 "
             "(minutes 2026-21). Owner: Esteban Ospina, Credit Director.")
RULES = [
    ("1. Scope", [
        "1.1 Libranza Empresarial is a payroll-deduction loan to a member employed by a private employer that has signed "
        "a deduction agreement with the cooperative. The libranza for public servants (Libranza Publica) is governed by "
        "its own rules and is not changed by these.",
    ]),
    ("2. Eligibility", [
        "2.1 The member has made savings contributions through the employer's payroll in each of the twelve months "
        "before the application.",
        "2.2 The instalment does not exceed 30 per cent of the member's net monthly pay in the employer's latest "
        "ordinary payroll, after the deductions already in force.",
        "2.3 The member is at most 65 years old at the loan's last instalment.",
    ]),
    ("3. Terms", [
        "3.1 Amounts from VMP 1,000,000 to VMP 40,000,000, over 12 to 60 months, at the monthly rate of the rate card in "
        "force on the disbursement date, as a fixed monthly instalment.",
        "3.2 The first instalment is deducted in the first payroll after disbursement. An employer that pays "
        "fortnightly deducts half of the instalment in each fortnight.",
    ]),
    ("4. The employer under the agreement", [
        "4.1 The employer uploads its payroll register for every pay period to the cooperative's intake portal within "
        "three business days of the pay date, ordinary, bonus and final-settlement payments included.",
        "4.2 The employer deducts each member's savings contribution and each instalment and transfers them to the "
        "collection account within five business days of the pay date, with the period in the transfer reference.",
        "4.3 The employer files a savings contribution schedule for each month through the portal.",
    ]),
    ("5. Default", [
        "5.1 A loan is in default from the month-end at which an instalment is 90 or more days past due, the definition "
        "the cooperative applies to every loan in its annual return.",
        "5.2 On default the cooperative may net the member's own savings against the loan's balance (Statutes, article "
        "41).",
    ]),
    ("6. Launch", [
        "6.1 The first Libranza Empresarial loans are disbursed together on 1 February 2027 to the applications approved "
        "in the launch disbursement file.",
    ]),
]
RULES_REVISIONS = ("Version 2.0 (20 November 2026): section 6 added; 2.1 to twelve months. Version 1.1 "
                   "(14 March 2025): section 4. Version 1.0 (30 January 2025).")

# ------------------------------------------------------------------------------------------ the supervisor return
SUPER_DATE = dt.date(2026, 3, 27)
SUPER_TITLE = "Annual return 2025: payroll-deduction portfolio"
SUPER_SUB = ("Coopservir to the Superintendencia de la Economia Solidaria de Valdemora, form F-12 summary. Filed 27 March "
             "2026. Figures at each 31 December.")


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def pdf_stamp(d, hh=10, mi=0):
    return ("D:%04d%02d%02d%02d%02d00-05'00'" % (d.year, d.month, d.day, hh, mi)).encode()


def _styles():
    return dict(
        org=ParagraphStyle("org", fontName="Carlito-Bold", fontSize=9.5, leading=12, textColor=ACCENT, spaceAfter=2),
        h1=ParagraphStyle("h1", fontName="LibSans-Bold", fontSize=15, leading=19, textColor=INK, spaceAfter=4),
        sub=ParagraphStyle("sub", fontName="Carlito", fontSize=9.2, leading=12, textColor=GREY, spaceAfter=12),
        h2=ParagraphStyle("h2", fontName="LibSans-Bold", fontSize=11, leading=14, textColor=ACCENT, spaceBefore=8,
                          spaceAfter=3),
        body=ParagraphStyle("body", fontName="Carlito", fontSize=10.8, leading=14.4, textColor=INK, spaceAfter=5),
        cell=ParagraphStyle("cell", fontName="Carlito", fontSize=9.4, leading=12, textColor=INK),
        small=ParagraphStyle("small", fontName="Carlito", fontSize=8.6, leading=11, textColor=GREY, spaceBefore=10))


def _footer(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Carlito", 7.8)
        canvas.setFillColor(GREY)
        canvas.drawString(20 * mm, 11 * mm, text)
        canvas.drawRightString(A4[0] - 20 * mm, 11 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _doc(path, title, author):
    return SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=20 * mm,
                             bottomMargin=20 * mm, title=title, author=author, subject="", creator="", producer="")


def _table(rows, widths, st):
    data = [[Paragraph(esc(str(c)), st["cell"]) for c in r] for r in rows]
    t = Table(data, colWidths=widths)
    t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, 0), 0.6, RULE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                           ("FONTNAME", (0, 0), (-1, 0), "Carlito-Bold")]))
    return t


def rules_text():
    out = [RULES_TITLE, RULES_SUB]
    for h, paras in RULES:
        out.append(h)
        out.extend(paras)
    out.append(RULES_REVISIONS)
    return out


def write_rules(out):
    path = os.path.join(out, C.F_RULES)
    st = _styles()
    doc = _doc(path, RULES_TITLE, "Credit Directorate")
    story = [Paragraph(esc(C.ORG + " | Credit Directorate"), st["org"]), Paragraph(esc(RULES_TITLE), st["h1"]),
             Paragraph(esc(RULES_SUB), st["sub"])]
    for h, paras in RULES:
        story.append(Paragraph(esc(h), st["h2"]))
        for p in paras:
            story.append(Paragraph(esc(p), st["body"]))
    story.append(Spacer(1, 4))
    story.append(Paragraph(esc(RULES_REVISIONS), st["small"]))
    foot = "Coopservir. Credit Directorate. Board approved product rules."
    doc.build(story, onFirstPage=_footer(foot), onLaterPages=_footer(foot))
    PP.normalize_pdf(path, pdf_stamp(RULES_DATE, 17, 5), pdf_stamp(RULES_DATE, 17, 5))


def supervisor_rows(P):
    """the libranza book at each 31 December 2021 to 2025, as the return reports it"""
    rows = []
    for y in range(2021, 2026):
        t = C.mi(y, 12)
        open_ = (P["disb"] <= t) & ((P["end_m"] < 0) | (P["end_m"] > t)) & ((P["dflt_m"] < 0) | (P["dflt_m"] > t))
        bal = 0.0
        for i in np.where(open_)[0]:
            bal += C.balance_after(P["principal"][i], P["rate"][i], int(P["term"][i]), int(t - P["disb"][i]))
        in_def = (P["dflt_m"] >= 0) & (P["dflt_m"] <= t) & (P["dflt_m"] > t - 24)
        dbal = 0.0
        for i in np.where(in_def)[0]:
            k = int(P["dflt_m"][i] - P["disb"][i]) - 3
            dbal += C.balance_after(P["principal"][i], P["rate"][i], int(P["term"][i]), max(0, k))
        tot = bal + dbal
        rows.append((str(y), "{:,}".format(int(open_.sum() + in_def.sum())), "{:,.0f}".format(tot / 1e6),
                     "%.2f%%" % (100 * dbal / tot), "{:,.0f}".format(dbal * 0.62 / 1e6)))
    return rows


def write_supervisor(out, P):
    path = os.path.join(out, C.F_SUPER)
    st = _styles()
    doc = _doc(path, SUPER_TITLE, "Finance")
    story = [Paragraph(esc(C.ORG + " | Finance"), st["org"]), Paragraph(esc(SUPER_TITLE), st["h1"]),
             Paragraph(esc(SUPER_SUB), st["sub"])]
    story.append(Paragraph(esc("Libranza Publica portfolio"), st["h2"]))
    story.append(_table([("31 December", "Loans", "Balance, VMP million", "90+ days past due, share of balance",
                          "Provision held, VMP million")] + supervisor_rows(P),
                        [24 * mm, 22 * mm, 36 * mm, 46 * mm, 38 * mm], st))
    story.append(Paragraph(esc("The balance includes loans in default for up to 24 months after default; loans written "
                               "off are excluded. Provisions follow the Superintendencia's schedule for payroll-deduction "
                               "consumer loans."), st["body"]))
    story.append(Paragraph(esc("Libranza Empresarial: no loans at 31 December 2025."), st["body"]))
    foot = "Coopservir. Form F-12 summary, 2025."
    doc.build(story, onFirstPage=_footer(foot), onLaterPages=_footer(foot))
    PP.normalize_pdf(path, pdf_stamp(SUPER_DATE, 11, 40), pdf_stamp(SUPER_DATE, 11, 40))


# ------------------------------------------------------------------------------------------------- Excel files
def write_agreements(out, E):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    path = os.path.join(out, C.F_AGREE)
    wb = Workbook()
    ws = wb.active
    ws.title = "agreements"
    head = ("employer_nit", "legal_name", "sector", "agreement_signed", "first_register_period", "city")
    fill = PatternFill("solid", fgColor="DCE7E2")
    for j, h in enumerate(head):
        c = ws.cell(row=1, column=j + 1, value=h)
        c.font = Font(bold=True)
        c.fill = fill
    rng = C.stream("agreements_xlsx")
    import world as W
    for i, e in enumerate(sorted(E, key=lambda e: e["start"])):
        y, m = C.ym(e["start"])
        signed = dt.date(y, m, 1) - dt.timedelta(days=int(rng.integers(5, 40)))
        city = W.PLACES[int(rng.integers(0, len(W.PLACES)))]
        vals = (e["nit"], e["name"], e["sector"], signed, "%04d-%02d" % (y, m), city)
        for j, v in enumerate(vals):
            c = ws.cell(row=i + 2, column=j + 1, value=v)
            if j == 3:
                c.number_format = "yyyy-mm-dd"
    for col, w in zip("ABCDEF", (16, 42, 22, 16, 20, 22)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(path)
    PP.normalize_ooxml(path, "Paola Ayala", "Paola Ayala", "2025-01-28T14:02:00Z", "2026-11-03T16:21:00Z",
                       excel=True, company=C.ORG)


FIELDS = (
    (C.F_LOANS, "loan_no", "text", "The loan's number in the core system. Loans closed before the core system went live in July 2021 keep their legacy number; legacy loans open at the move carry a new number (see the crosswalk)."),
    (C.F_LOANS, "member_id", "integer", "The member's national id number."),
    (C.F_LOANS, "payroll_unit", "text", "The Tesoreria payroll unit that pays the member's salary and takes the deduction."),
    (C.F_LOANS, "unit_type", "text", "national, territorial, health or education, from the Tesoreria's unit list."),
    (C.F_LOANS, "contract_type", "text", "The member's appointment at origination: CARRERA, LIBRE NOMBRAMIENTO or TEMPORAL in the core system; P (permanent) or T (temporary) as stored by the legacy system."),
    (C.F_LOANS, "hire_date", "date", "First day of the month the member entered the payroll unit's service."),
    (C.F_LOANS, "birth_year", "integer", "The member's year of birth."),
    (C.F_LOANS, "net_pay_at_origination", "number", "Net monthly pay on the payroll certificate at origination, in pesos; legacy-system loans in thousands of pesos."),
    (C.F_LOANS, "disbursement_date", "date", "Date the loan was disbursed."),
    (C.F_LOANS, "principal", "number", "Amount lent, in pesos; legacy-system loans in thousands of pesos."),
    (C.F_LOANS, "term_months", "integer", "Number of monthly instalments."),
    (C.F_LOANS, "monthly_rate", "number", "Monthly interest rate."),
    (C.F_LOANS, "instalment", "number", "Fixed monthly instalment, in pesos; legacy-system loans in thousands of pesos."),
    (C.F_LOANS, "collection_channel", "text", "How the loan was being collected on 31 December 2026: NOMINA (payroll deduction), DEBITO DIRECTO (debit to the member's account), COBRANZA (in collections after default) or CERRADO."),
    (C.F_LOANS, "closed_date", "date", "Date the loan closed, for CERRADO loans."),
    (C.F_LOANS, "closure_reason", "text", "VENCIMIENTO (paid to term) or COMPRA DE CARTERA (paid off by another lender)."),
    (C.F_LOANS, "origin_system", "text", "LEGACY for loans disbursed in the legacy system (to June 2021), CORE after."),
    (C.F_XWALK, "legacy_loan_no", "text", "Legacy loan number."),
    (C.F_XWALK, "core_loan_no", "text", "Number given in the core system to a legacy loan open at the move."),
    (C.F_XWALK, "migrated_on", "date", "Migration date."),
    (C.F_RETURNS, "period", "text", "Payroll month of the return: YYYYMM to March 2023, YYYY-MM from April 2023."),
    (C.F_RETURNS, "payroll_unit", "text", "The Tesoreria payroll unit."),
    (C.F_RETURNS, "national_id", "integer", "The member's national id number."),
    (C.F_RETURNS, "loan_reference", "text", "The loan number the deduction order carries (legacy numbers to June 2021)."),
    (C.F_RETURNS, "amount", "number", "Instalment not deducted, in pesos (legacy loans: as ordered, in pesos)."),
    (C.F_RETURNS, "reason", "text", "Why the deduction was not applied: 01 or SIN SALARIO, no salary paid to the member in the month; 02 or SIN CUPO, salary paid but insufficient margin after priority deductions. One row per deduction not applied; deductions applied are not returned."),
    (C.F_DEFAULTS, "loan_reference", "text", "The loan's number at the default date (legacy numbers to June 2021)."),
    (C.F_DEFAULTS, "member_id", "integer", "The member's national id number."),
    (C.F_DEFAULTS, "default_date", "date", "The month-end at which an instalment first stood 90 or more days past due."),
    (C.F_DEFAULTS, "balance_at_default", "number", "Principal balance owed at the default date, in pesos."),
    (C.F_DEFAULTS, "cure_date", "date", "Date the arrears were cleared, if they were."),
    (C.F_RECEIPTS, "receipt_id", "text", "Receipt number in the collections system."),
    (C.F_RECEIPTS, "receipt_date", "date", "Date the payment was received."),
    (C.F_RECEIPTS, "loan_reference", "text", "The defaulted loan the payment was applied to."),
    (C.F_RECEIPTS, "member_id", "integer", "The borrowing member's national id number."),
    (C.F_RECEIPTS, "channel", "text", "TRANSFERENCIA (transfer into the collection account), CAJA (cash at a branch) or REVERSO (a receipt reversed because the payment was returned)."),
    (C.F_RECEIPTS, "originating_account", "text", "For transfers, the paying account as bank code and account number."),
    (C.F_RECEIPTS, "amount", "number", "Amount applied, in pesos; reversals negative."),
    (C.F_SAVINGS, "savings_account", "text", "The member's savings account at the cooperative, bank code 0731."),
    (C.F_SAVINGS, "member_id", "integer", "The account holder's national id number."),
    (C.F_SAVINGS, "movement_date", "date", "Date of the movement."),
    (C.F_SAVINGS, "movement", "text", "SALDO INICIAL (balance on 1 January 2023), APORTE (contribution), INTERES (annual interest) or RETIRO (withdrawal or transfer out)."),
    (C.F_SAVINGS, "amount", "number", "Amount in pesos, withdrawals negative."),
    (C.F_SAVINGS, "balance", "number", "Balance after the movement."),
    (C.F_REGISTERS, "employer_nit", "text", "The employer's tax number."),
    (C.F_REGISTERS, "register_no", "text", "The employer's register number for the upload."),
    (C.F_REGISTERS, "period_start", "date", "First day of the pay period."),
    (C.F_REGISTERS, "period_end", "date", "Last day of the pay period."),
    (C.F_REGISTERS, "pay_date", "date", "Date the payment was made."),
    (C.F_REGISTERS, "payment_type", "text", "NOMINA (ordinary pay), PRIMA (the June and December service bonus) or LIQUIDACION (final settlement at the end of employment)."),
    (C.F_REGISTERS, "employee_id", "text", "The employee's national id as the intake portal stored it."),
    (C.F_REGISTERS, "hire_date", "date", "First day of the month the employee was hired."),
    (C.F_REGISTERS, "base_salary", "number", "Monthly base salary, in pesos."),
    (C.F_REGISTERS, "gross_pay", "number", "Gross pay for the period, in pesos."),
    (C.F_REGISTERS, "deductions", "number", "All deductions for the period, in pesos."),
    (C.F_REGISTERS, "net_pay", "number", "Net pay for the period, in pesos."),
    (C.F_LAUNCH, "loan_no", "text", "Libranza Empresarial loan number."),
    (C.F_LAUNCH, "member_id", "integer", "The member's national id number."),
    (C.F_LAUNCH, "employer_nit", "text", "The employer's tax number."),
    (C.F_LAUNCH, "hire_date", "date", "From the employer's employment certificate."),
    (C.F_LAUNCH, "member_since", "date", "Month the member's first contribution was received."),
    (C.F_LAUNCH, "certified_net_pay", "number", "Net monthly pay certified by the employer for the application, in pesos."),
    (C.F_LAUNCH, "principal", "number", "Amount to be disbursed, in pesos."),
    (C.F_LAUNCH, "term_months", "integer", "Number of monthly instalments."),
    (C.F_LAUNCH, "monthly_rate", "number", "Monthly interest rate."),
    (C.F_LAUNCH, "instalment", "number", "Fixed monthly instalment, in pesos."),
    (C.F_LAUNCH, "disbursement_date", "date", "Disbursement date."),
    (C.F_LAUNCH, "first_deduction_period", "text", "Payroll month of the first deduction."),
    (C.F_SCHED, "schedule_no", "text", "Number of the savings contribution schedule."),
    (C.F_SCHED, "employer_nit", "text", "The employer's tax number."),
    (C.F_SCHED, "period", "text", "Month of the contributions."),
    (C.F_SCHED, "filed_on", "date", "Date the schedule was filed through the portal."),
    (C.F_SCHED, "member_id", "integer", "The member's national id number."),
    (C.F_SCHED, "contribution", "number", "Savings contribution deducted, in pesos."),
    (C.F_STMT, "posting_date", "date", "Posting date on the collection account."),
    (C.F_STMT, "description", "text", "The bank's transaction description."),
    (C.F_STMT, "payer_id", "text", "Tax number or national id of the payer, as the paying bank sent it."),
    (C.F_STMT, "payer_name", "text", "Payer's name, as the paying bank sent it."),
    (C.F_STMT, "reference", "text", "Reference sent with the payment."),
    (C.F_STMT, "credit", "number", "Credit, in pesos."),
    (C.F_STMT, "debit", "number", "Debit, in pesos."),
    (C.F_STMT, "balance", "number", "Running balance from 1 January 2025, in pesos."),
    (C.F_LABOUR, "trimestre", "text", "Quarter."),
    (C.F_LABOUR, "rama_actividad", "text", "Branch of economic activity."),
    (C.F_LABOUR, "tasa_desempleo_pct", "number", "Unemployment rate, per cent."),
    (C.F_LABOUR, "ocupados_miles", "number", "Employed persons, thousands."),
    (C.F_AGREE, "employer_nit", "text", "The employer's tax number."),
    (C.F_AGREE, "legal_name", "text", "Registered name."),
    (C.F_AGREE, "sector", "text", "The cooperative's sector for the employer."),
    (C.F_AGREE, "agreement_signed", "date", "Date the deduction agreement was signed."),
    (C.F_AGREE, "first_register_period", "text", "First month for which the employer uploaded its payroll register."),
    (C.F_AGREE, "city", "text", "City of the employer's main site."),
)


def write_dictionary(out):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    path = os.path.join(out, C.F_DICT)
    wb = Workbook()
    ws = wb.active
    ws.title = "fields"
    head = ("file", "field", "type", "description")
    fill = PatternFill("solid", fgColor="DCE7E2")
    for j, h in enumerate(head):
        c = ws.cell(row=1, column=j + 1, value=h)
        c.font = Font(bold=True)
        c.fill = fill
    for i, row in enumerate(FIELDS):
        for j, v in enumerate(row):
            c = ws.cell(row=i + 2, column=j + 1, value=v)
            if j == 3:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCD", (44, 26, 10, 100)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(path)
    PP.normalize_ooxml(path, "Paola Ayala", "Paola Ayala", "2025-02-11T15:40:00Z", "2027-01-21T18:12:00Z",
                       excel=True, company=C.ORG)
