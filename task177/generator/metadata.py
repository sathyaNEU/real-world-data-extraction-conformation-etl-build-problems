"""Write task177/metadata.json from the built pack: what every file is and where it comes from, what is constructed,
the distractors, the input gates."""
from __future__ import annotations

import csv
import json
import os

from openpyxl import load_workbook

import common as C

ORIGIN = "Constructed for this task (seeded, deterministic): "
SOURCES = {
    C.F_LOANS: ("Coopservir (fictional) core banking system, Libranza Publica loan book extract at 31 December 2026, "
                "loans disbursed January 2017 to December 2026", "2027-01-22"),
    C.F_XWALK: ("Coopservir core system migration, legacy to core loan numbers for the loans open at the July 2021 "
                "move", "2021-07-01"),
    C.F_RETURNS: ("Tesoreria General de Valdemora (fictional) monthly returns of payroll deductions not applied, as "
                  "loaded by Coopservir", "2027-01-22"),
    C.F_DEFAULTS: ("Coopservir credit risk, register of Libranza Publica defaults (90 days past due)", "2027-01-22"),
    C.F_REGISTERS: ("Coopservir intake portal, payroll registers uploaded by the 38 agreement employers", "2027-01-22"),
    C.F_AGREE: ("Coopservir commercial directorate, register of Libranza Empresarial deduction agreements",
                "2026-11-03"),
    C.F_LAUNCH: ("Coopservir credit directorate, launch disbursement file of approved Libranza Empresarial "
                 "applications", "2027-01-22"),
    C.F_SCHED: ("Coopservir intake portal, savings contribution schedules filed by the agreement employers",
                "2027-01-22"),
    C.F_STMT: ("Coopservir treasury, collection account statement export", "2027-01-22"),
    C.F_RECEIPTS: ("Coopservir collections system, receipts applied to defaulted Libranza Publica loans", "2027-01-22"),
    C.F_SAVINGS: ("Coopservir savings system, ledger extract for members with loans in collections and other members "
                  "whose accounts paid them", "2027-01-22"),
    C.F_RULES: ("Coopservir credit directorate, Libranza Empresarial product rules, version 2.0", "2026-11-20"),
    C.F_DICT: ("Coopservir finance, field descriptions for the pull", "2027-01-21"),
    C.F_SUPER: ("Coopservir finance, 2025 annual return to the supervisor (fictional), form F-12 summary", "2026-03-27"),
    C.F_LABOUR: ("Instituto Nacional de Estadistica de Valdemora (fictional), quarterly labour market bulletin, "
                 "unemployment and employment by branch", "2026-11-14"),
}
TITLE = ("How many of a public servants' savings and credit cooperative's first 2,400 payroll-deduction loans to "
         "members employed by private firms will go into default within the first year, forecast from ten years of "
         "payroll loans to public servants and the private employers' own payroll registers")


def _rows(path):
    ext = path.rsplit(".", 1)[-1]
    if ext == "csv":
        with open(path, encoding="utf-8", newline="") as fh:
            return sum(1 for _ in csv.reader(fh)) - 1
    if ext == "xlsx":
        wb = load_workbook(path, read_only=True)
        return sum(max(0, ws.max_row - 1) for ws in wb.worksheets)
    return None


def write(task_dir):
    target = os.path.join(task_dir, "target")
    files = []
    big, big_rows = None, 0
    for f in sorted(os.listdir(target)):
        p = os.path.join(target, f)
        rows = _rows(p)
        if rows and rows > big_rows:
            big, big_rows = f, rows
        src, date = SOURCES[f]
        files.append(dict(path=f, format=f.rsplit(".", 1)[-1], bytes=os.path.getsize(p), rows=rows,
                          source=ORIGIN + src, date=date, license="CC0-1.0 (original work, constructed for this task)"))
    distractors = [
        dict(file=C.F_SUPER, wrong_basis=False,
             why_relevant="the cooperative's own reported default share of its payroll-deduction book by year, which "
                          "looks like the benchmark rate for a new payroll-deduction product",
             why_unused="it reports the public book's balance in default at each year end, a stock on a balance basis, "
                        "not a first-year default count; no figure reads it",
             ruled_out_by="its own basis (balance 90+ days past due at 31 December, defaults kept for 24 months) and "
                          "the product rules' default definition"),
        dict(file=C.F_LABOUR, wrong_basis=False,
             why_relevant="unemployment and employment by branch of activity, which looks like a way to adjust a "
                          "default rate for the private sectors' labour markets",
             why_unused="national rates by branch say nothing about the 38 employers' own staff; their payroll "
                        "registers do; no figure reads it",
             ruled_out_by="the employers' payroll registers, which record their own staff month by month"),
    ]
    meta = dict(
        task="task177", title=TITLE,
        domain="Banking, Credit & Financial Risk", client_domain="banking-credit-financial-risk",
        subdomain="credit scoring and PD modelling (the first-year default of payroll-deduction loans to a new "
                  "population of borrowers whose employers' staff leave far more often than the record's)",
        objective="Forecasting & Predictive Modeling",
        prompt_shape="one forecast count to the nearest 50 loans and each employer sector's share in default to the "
                     "nearest 5 per cent; the savings deducted for 2026 and not received by year end by sector; the "
                     "recoveries on 2023 and 2024 public defaults by unit type",
        as_of="2027-01-22",
        deliverables=["libranza_empresarial_default_forecast.xlsx", "launch_default_note.docx"],
        distractor_files=[d["file"] for d in distractors],
        distractors=distractors,
        stated_inputs=["2,400", "1 February 2027", "31 January 2028"],
        field_terms=[],
        input_gates=dict(files=len(files), formats=sorted(set(x["format"] for x in files)), largest_file=big,
                         largest_file_rows=big_rows),
        sources=[dict(files="all files under target/",
                      origin="constructed for this task by task177/generator/build.py (seeded, deterministic) around a "
                             "fictional public servants' savings and credit cooperative (Coopservir) in a fictional "
                             "country (Republica de Valdemora), its payroll-deduction loans, the fictional national "
                             "treasury's returns and 38 fictional private employers. Members and employees are "
                             "identified by number only. No real organisation's records or person is depicted.",
                      date="2027-01-22", license="CC0-1.0 (original work)")],
        created="2026-10-23", generator="task177/generator/build.py", files=files)
    with open(os.path.join(task_dir, "metadata.json"), "w") as fh:
        json.dump(meta, fh, indent=1)
        fh.write("\n")
    return meta
