"""task178 pack: writes every file under target/ from world.py and docs.py."""
from __future__ import annotations

import datetime as dt
import math
import os

import numpy as np
import openpyxl
from openpyxl.styles import Font

import common as C
import docs as D
import postprocess as PP
import world as W


def wb_write(path, sheets, stamp, widths=None, formats=None, company=C.COMPANY):
    wb = openpyxl.Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        first = False
        for r in rows:
            ws.append(list(r))
        if rows:
            for c in ws[1]:
                c.font = Font(bold=True)
        for col, fmt in (formats or {}).get(name, {}).items():
            for cell in ws[col][1:]:
                cell.number_format = fmt
        for k, w in (widths or {}).get(name, {}).items():
            ws.column_dimensions[k].width = w
    wb.save(path)
    creator, last_by, created, modified = stamp
    PP.normalize_ooxml(path, creator, last_by, created, modified, excel=True, company=company)
    return sum(len(r) - 1 for r in sheets.values())


_SHUT = C.closed_days()


def next_day(d, days_set=None):
    """the next day the platform opens and both currencies settle"""
    x = d + dt.timedelta(days=1)
    while x.weekday() >= 5 or x in _SHUT:
        x += dt.timedelta(days=1)
    return x


def mid_at(i, h, frac):
    R, on, mids = W.path()
    return mids[i, h] * math.exp(R[i, h] / 100.0 * frac)


# ------------------------------------------------------------------------------------------------ the Bank's operations
def write_ops(out):
    rng = C.stream("pack_ops")
    days = W.days()
    dset = set(days)
    _, _, tickets = W.bank()
    rows = []
    for (i, h, m, s, amt, dealer, cancel) in tickets:
        t_loc = dt.datetime.combine(days[i], dt.time(C.HOURS[h], m, s))
        price = round(mid_at(i, h, (m * 60 + s) / 3600.0) * (1 + (0.00004 if amt > 0 else -0.00004)), 2)
        vd = next_day(days[i], dset)
        op = "VENTA" if amt > 0 else "COMPRA"
        if cancel:
            wrong = abs(amt) + float(rng.choice([-2.0, -1.0, 1.0, 2.0, 5.0])) if abs(amt) > 2.5 else abs(amt) + 1.0
            rows.append([t_loc - dt.timedelta(seconds=int(rng.integers(40, 170))), vd, op, wrong, dealer, price, "ANULADA"])
        rows.append([t_loc, vd, op, abs(amt), dealer, price, "VIGENTE"])
    rows.sort(key=lambda r: r[0])
    out_rows, seq = [], {}
    for r in rows:
        y = r[0].year
        seq[y] = seq.get(y, 0) + 1
        out_rows.append(["OC%d-%05d" % (y, seq[y]), r[0].strftime("%Y-%m-%d %H:%M:%S"), r[1].isoformat(), r[2],
                         "%d" % round(r[3] * 1000), r[4], "%.2f" % r[5], r[6]])
    return C.write_csv(os.path.join(out, C.F_OPS), ["operation_id", "dealt_at", "value_date", "operation",
                                                     "amount_usd_k", "counterparty_dealer", "price_tqp", "status"],
                       out_rows)


# ------------------------------------------------------------------------------------------------ the dealers' reports
def write_flows(out):
    rng = C.stream("pack_flows")
    days = W.days()
    flows = W.dealer_flows()
    fam_raw, jul = W.fam_legacy()
    codes = [d[0] for d in W.DEALERS]
    fam = codes.index("FTV")
    rows = []
    for i, d in enumerate(days):
        for k, code in enumerate(codes):
            if code == "FTV" and d >= C.MERGER_DATE and not jul[i]:
                continue
            base = dt.datetime.combine(d, dt.time(16, 0)) + dt.timedelta(seconds=int(rng.integers(240, 6900)))
            for h, hh in enumerate(C.HOURS):
                v = fam_raw[i, h] if (code == "FTV" and d >= C.MERGER_DATE) else flows[k, i, h]
                lab = "%02d-%02d" % (hh, hh + 1)
                rec = base + dt.timedelta(seconds=h * int(rng.integers(0, 3)))
                if rng.random() < 0.0075:
                    kind = rng.random()
                    if kind < 0.35:
                        first = -v
                    elif kind < 0.6:
                        first = flows[k, i, h - 1] if h > 0 else round(v * 0.5)
                    else:
                        first = round(v * rng.uniform(0.3, 0.8))
                    rows.append([d, lab, code, int(first), rec])
                    late = rec + dt.timedelta(minutes=int(rng.integers(35, 900)))
                    rows.append([d, lab, code, int(v), late])
                else:
                    rows.append([d, lab, code, int(v), rec])
    rows.sort(key=lambda r: (r[0], r[1], r[2], r[4]))
    out_rows = [[r[0].isoformat(), r[1], r[2], r[3], r[4].strftime("%Y-%m-%d %H:%M:%S")] for r in rows]
    return C.write_csv(os.path.join(out, C.F_FLOWS), ["report_date", "hour", "dealer_code", "net_customer_usd_k",
                                                       "received_at"], out_rows)


# ------------------------------------------------------------------------------------------------ the hourly mids
def write_mids(out):
    days = W.days()
    _, _, mids = W.path()
    lost = set(W.OUTAGES)
    rows = []
    for i, d in enumerate(days):
        for j in range(len(C.HOURS) + 1):
            hh = C.OPEN_H + j
            if (d, hh) in lost:
                continue
            rows.append([d.isoformat(), "%02d:00" % hh, "%.2f" % mids[i, j]])
    return C.write_csv(os.path.join(out, C.F_MID), ["session_date", "hour_mark", "mid_tqp_per_usd"], rows)


def write_dealers(out):
    rows = []
    for code, name, _ in W.DEALERS:
        frm, to, note = "2024-01-02", "", ""
        if code == "FTV":
            to, note = "2025-06-30", "Merged into Banco del Salitral (BSL) effective 1 July 2025"
        if code == "BSL":
            note = "Reports include the former Financiera Tres Valles customers from 1 July 2025"
        rows.append([code, name, frm, to, note])
    return C.write_csv(os.path.join(out, C.F_DEALERS), ["dealer_code", "institution", "of_record_from", "of_record_to",
                                                         "note"], rows)


def write_calendar(out):
    shut = C.closed_days()
    rows, d = [], dt.date(2024, 1, 1)
    while d <= dt.date(2026, 12, 31):
        wd = d.strftime("%a")
        if d.weekday() >= 5:
            rows.append([d.isoformat(), wd, "CLOSED", "weekend"])
        elif d in shut:
            rows.append([d.isoformat(), wd, "CLOSED", "; ".join(shut[d])])
        elif d > C.LAST_DAY:
            rows.append([d.isoformat(), wd, "OPEN", "after the release period"])
        else:
            rows.append([d.isoformat(), wd, "OPEN", ""])
        d += dt.timedelta(days=1)
    return C.write_csv(os.path.join(out, C.F_CAL), ["date", "weekday", "platform", "reason"], rows)


# ------------------------------------------------------------------------------------------------ the TRM and the monthly statistics
def trm_series():
    rng = C.stream("trm")
    days = W.days()
    _, _, mids = W.path()
    w = np.array([1.6, 1.3, 1.1, 0.9, 0.8, 0.9, 1.1, 1.5])
    out = {}
    for i, d in enumerate(days):
        seg = (mids[i, :-1] + mids[i, 1:]) / 2
        out[d] = round(float((seg * w).sum() / w.sum()) + rng.normal(0, 0.6), 2)
    return out


def write_trm(out):
    days = W.days()
    vw = trm_series()
    rows, d, cur = [], dt.date(2024, 1, 3), vw[days[0]]
    last_session = {}
    idx = 0
    while d <= dt.date(2026, 12, 30):
        while idx < len(days) and days[idx] < d:
            cur = vw[days[idx]]
            idx += 1
        rows.append([d.isoformat(), "%.2f" % cur])
        d += dt.timedelta(days=1)
    return C.write_csv(os.path.join(out, C.F_TRM), ["fecha_vigencia", "trm_tqp_per_usd"], rows)


def write_stats(out):
    rng = C.stream("stats")
    days = W.days()
    dset = set(days)
    flows = W.dealer_flows().sum(axis=0) / 1000.0
    _, _, tickets = W.bank()
    by_value = {}
    for (i, h, m, s, amt, dealer, cancel) in tickets:
        vd = next_day(days[i], dset)
        key = (vd.year, vd.month)
        by_value[key] = by_value.get(key, 0.0) + amt
    months = [(y, m) for y in (2024, 2025, 2026) for m in range(1, 13) if (y, m) <= (2026, 11)]
    rows = [["Month", "Interbank turnover (USD mn)", "Dealers' net customer purchases (USD mn)",
             "BCT net spot sales, value date (USD mn)", "International reserves, end of month (USD mn)"]]
    res = 19850.0
    for (y, m) in months:
        idx = [i for i, d in enumerate(days) if (d.year, d.month) == (y, m)]
        cust = float(flows[idx].sum()) if idx else 0.0
        turn = sum(1180.0 * W.day_scale()[i] * rng.uniform(0.85, 1.15) for i in idx)
        sales = by_value.get((y, m), 0.0)
        res = res - sales + rng.normal(95.0, 140.0)
        rows.append(["%04d-%02d" % (y, m), round(turn, 0), round(cust, 1), round(sales, 1), round(res, 0)])
    return wb_write(os.path.join(out, C.F_STATS), {"Mercado cambiario": rows},
                    ("Departamento de Estadísticas", "Departamento de Estadísticas", "2024-02-06T15:12:40Z",
                     "2026-12-07T14:03:22Z"), widths={"Mercado cambiario": {"A": 10, "B": 26, "C": 36, "D": 34, "E": 38}},
                    company=C.BANK)


# ------------------------------------------------------------------------------------------------ the desk's records
COUNTERPARTIES = [d[0] for d in W.DEALERS if d[0] not in ("BCO", "FTV")]


def write_confirmations(out):
    rng = C.stream("pack_conf")
    rows = []
    seq = 104300
    for o in W.desk_orders():
        utc = C.local_to_utc(o["local"])
        stamp = utc.strftime("%Y-%m-%dT%H:%M:%S.") + "%03dZ" % (utc.microsecond // 1000)
        for px, q in o["fills"]:
            cp = COUNTERPARTIES[int(rng.integers(0, len(COUNTERPARTIES)))]
            if rng.random() < 0.012:
                seq += int(rng.integers(1, 4))
                bad = round(px + (1 if rng.random() < 0.5 else -1) * float(rng.choice([10.0, 1.0, 100.0])), 2)
                rows.append(["PL26-%07d" % seq, stamp, o["trader"], o["side"], int(q), "%.2f" % bad,
                             "%.2f" % o["bid"], "%.2f" % o["ask"], cp, "CANCELLED"])
            seq += int(rng.integers(1, 4))
            rows.append(["PL26-%07d" % seq, stamp, o["trader"], o["side"], int(q), "%.2f" % px, "%.2f" % o["bid"],
                         "%.2f" % o["ask"], cp, "CONFIRMED"])
    return C.write_csv(os.path.join(out, C.F_CONF), ["confirmation_id", "trade_time_utc", "trader", "side",
                                                      "usd_amount_k", "price", "bid_at_arrival", "ask_at_arrival",
                                                      "counterparty", "status"], rows)


def write_clients(out):
    rng = C.stream("pack_clients")
    branches = ("Ilvara Centro", "Puerto Quindal", "Santa Ofelia", "Valle Lindo", "Corporate Banking")
    rows = []
    for c in sorted(W.clients(), key=lambda c: c["no"]):
        since = dt.date(2009, 1, 1) + dt.timedelta(days=int(rng.integers(0, 5600)))
        br = "Treasury" if c["segment"] == "TESORERIA" else branches[int(rng.integers(0, len(branches)))]
        rows.append([c["no"], c["segment"], c["sector"], br, since.isoformat()])
    return C.write_csv(os.path.join(out, C.F_CLIENTS), ["client_no", "segment", "sector", "relationship_unit",
                                                         "client_since"], rows)


def deal_status(d):
    if d["cancelled"]:
        return "ANULADA"
    return "LIQUIDADA" if d["mat"] < C.AS_OF else "VIGENTE"


def write_forwards(out):
    rows = [["Deal No.", "Client No.", "Trade date", "Maturity date", "Product", "Window opens", "Client USD",
             "Pair", "Notional", "Forward rate", "Status"]]
    for d in W.forward_book():
        rows.append([d["no"], d["client"], d["td"], d["mat"], d["product"], d["window_open"], d["side"], d["pair"],
                     d["amount"], d["rate"], deal_status(d)])
    return wb_write(os.path.join(out, C.F_FWD), {"Deals": rows},
                    ("Gabriel Arias", "Luisa Molano", "2025-01-06T13:41:09Z", "2026-12-30T00:48:37Z"),
                    widths={"Deals": {"A": 10, "B": 11, "C": 12, "D": 13, "E": 11, "F": 13, "G": 11, "H": 10, "I": 14,
                                      "J": 12, "K": 11}},
                    formats={"Deals": {"C": "yyyy-mm-dd", "D": "yyyy-mm-dd", "F": "yyyy-mm-dd", "I": "#,##0",
                                       "J": "#,##0.00"}})


def write_docs(out):
    for f, doc in D.DOCS.items():
        D.build_pdf(os.path.join(out, f), doc)


def build_pack(out):
    os.makedirs(out, exist_ok=True)
    n = {}
    n[C.F_OPS] = write_ops(out)
    n[C.F_FLOWS] = write_flows(out)
    n[C.F_MID] = write_mids(out)
    n[C.F_DEALERS] = write_dealers(out)
    n[C.F_CAL] = write_calendar(out)
    n[C.F_TRM] = write_trm(out)
    n[C.F_STATS] = write_stats(out)
    n[C.F_CONF] = write_confirmations(out)
    n[C.F_CLIENTS] = write_clients(out)
    n[C.F_FWD] = write_forwards(out)
    write_docs(out)
    return n
