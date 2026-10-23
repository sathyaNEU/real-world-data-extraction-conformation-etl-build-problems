"""Write task178/metadata.json from the built pack: what every file is, where it comes from, the distractors and the
input gates."""
from __future__ import annotations

import json
import os

from openpyxl import load_workbook

import common as C

ORIGIN = "Constructed for this task (seeded, deterministic): "
SOURCES = {
    C.F_OPS: ("Banco Central de Tamarquia (fictional), Market Operations, spot operations with reporting dealers, 2 January 2024 to 29 December 2026", "2026-12-30"),
    C.F_FLOWS: ("Banco Central de Tamarquia, reporting dealers' hourly net customer purchases of US dollars, every report received", "2026-12-30"),
    C.F_MID: ("Banco Central de Tamarquia, the interbank platform's mid rate at each hour mark of the session", "2026-12-30"),
    C.F_DEALERS: ("Banco Central de Tamarquia, register of reporting dealers", "2026-12-30"),
    C.F_CAL: ("Banco Central de Tamarquia, platform calendar 2024 to 2026", "2026-12-30"),
    C.F_ANNOUNCE: ("Banco Central de Tamarquia, Press release 41/2026 (English version)", "2026-10-01"),
    C.F_NOTE: ("Banco Central de Tamarquia, Market Operations, release note for the hourly data", "2026-12-30"),
    C.F_POLICY: ("Banco Cordillera (fictional), Treasury, FX options pricing policy v2.1", "2026-12-15"),
    C.F_CONF: ("Interbank platform trade confirmations for Banco Cordillera, 2026", "2026-12-30"),
    C.F_FWD: ("Banco Cordillera, Treasury deal system, client FX forward deals traded 2025 to 2026", "2026-12-30"),
    C.F_CLIENTS: ("Banco Cordillera, corporate client register (FX segment)", "2026-12-30"),
    C.F_TRM: ("Banco Central de Tamarquia, representative market rate (TRM) by date of validity", "2026-12-30"),
    C.F_STATS: ("Banco Central de Tamarquia, monthly foreign exchange market statistics, January 2024 to November 2026", "2026-12-07"),
}
TITLE = ("The volatility a domestic bank's FX options desk prices its first one-year dollar options at for 2027, once "
         "the central bank's spot operations against the market's hourly moves end")


def _rows(path):
    ext = path.rsplit(".", 1)[-1]
    if ext == "csv":
        with open(path, encoding="utf-8") as fh:
            return sum(1 for _ in fh) - 1
    if ext == "xlsx":
        wb = load_workbook(path, read_only=True)
        return sum(max(0, ws.max_row - 1) for ws in wb.worksheets)
    return None


def write(task_dir):
    target = os.path.join(task_dir, "target")
    files = []
    big, big_rows = None, 0
    for f in sorted(os.listdir(target)):
        p = os.path.join(target, f)
        rows = _rows(p)
        if rows and rows > big_rows:
            big, big_rows = f, rows
        src, date = SOURCES[f]
        files.append(dict(path=f, format=f.rsplit(".", 1)[-1], bytes=os.path.getsize(p), rows=rows,
                          source=ORIGIN + src, date=date, license="CC0-1.0 (original work, constructed for this task)"))
    distractors = [
        dict(file=C.F_TRM, wrong_basis=False,
             why_relevant="the country's official dollar rate, published by the central bank every day of the period",
             why_unused="the options settle on the platform's 16:00 closing rate and the policy measures the closing rate's "
                        "daily changes; no figure reads the TRM",
             ruled_out_by="the pricing policy, sections 1 and 2 (cash settlement against the 16:00 closing rate)"),
        dict(file=C.F_STATS, wrong_basis=False,
             why_relevant="the central bank's monthly market statistics, with its net spot sales and the dealers' net "
                          "customer purchases by month",
             why_unused="monthly totals by value date cannot separate the hours the Bank traded in; every figure reads the "
                        "hourly record",
             ruled_out_by="the release note, section 5 (the Bank's sales there are by value date and by month)"),
    ]
    meta = dict(
        task="task178", title=TITLE,
        domain="Quantitative Finance", client_domain="Quantitative Finance",
        subdomain="volatility and derivatives pricing (the free-float volatility of an exchange rate whose record was "
                  "produced under a central bank's leaning operations, for pricing one-year FX options)",
        objective="Forecasting & Predictive Modeling",
        prompt_shape="one dial (the desk's whole-per-cent volatility grid, settings 8 to 13 per cent) against the "
                     "expected 2027 volatility; a 2026 hedging-cost table by quarter and order size; 2027 client forward "
                     "maturities by quarter and segment; a chart of the grid against the expected volatility",
        as_of=C.AS_OF.isoformat(),
        deliverables=[C.G_BOOK, C.G_CHART],
        distractor_files=[d["file"] for d in distractors],
        distractors=distractors,
        stated_inputs=["252 sessions", "16:00 closing rate", "2024 to 2026", "31 December 2026", "UTC-5"],
        field_terms=["realised volatility", "order flow", "price impact", "closing rate", "window forward", "arrival mid"],
        input_gates=dict(files=len(files), formats=sorted(set(x["format"] for x in files)), largest_file=big,
                         largest_file_rows=big_rows),
        sources=[dict(files="all files under target/",
                      origin="constructed for this task by task178/generator/build.py (seeded, deterministic) around a "
                             "fictional Andean republic (Tamarquia), its central bank, its currency, its reporting dealers "
                             "and a fictional domestic bank (Banco Cordillera). No real organisation's records or person "
                             "is depicted.",
                      date=C.AS_OF.isoformat(), license="CC0-1.0 (original work)")],
        created="2026-10-23", generator="task178/generator/build.py", files=files)
    with open(os.path.join(task_dir, "metadata.json"), "w") as fh:
        json.dump(meta, fh, indent=1, ensure_ascii=False)
        fh.write("\n")
    return meta
