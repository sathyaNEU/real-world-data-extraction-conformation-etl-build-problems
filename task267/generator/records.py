"""The program's exports, written from the clean truth: the screening tablets' encounters, the laboratory's OGTT
results, the candidate rules, the season's close-out report, the event calendar, the analysers' run log and asset
register, the laboratory's invoice lines, the post-event feedback and the region's chronic disease profile."""
import csv
import datetime as dt

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import askworld as AW
import common as K

R_SEED = 26717


def rng_for(*keys):
    return np.random.default_rng([R_SEED] + [int(k) for k in keys])


def _head(ws, header, widths, fill="2F5D50"):
    ws.append(header)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=fill)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A%d" % (ws.max_row + 1)


def _csv(path, header, rows, delimiter=","):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, delimiter=delimiter, lineterminator="\r\n")
        w.writerow(header)
        w.writerows(rows)
    return len(rows)


# ---------------------------------------------------------------------------------------------- encounters
def write_encounters(path, T):
    """One row per screening as the tablets synced them; a tablet that lost its connection at an event re-sent the
    event's rows when it reconnected, so a few rows arrive twice with the same encounter ID."""
    rng = rng_for(1)
    resend = {e["eid"] for e in T["E"] if rng.random() < 0.035}
    rows = []
    for x in T["enc"]:
        r = [x["enc"], x["a"]["pid"], x["event"]["eid"], x["time"].strftime("%Y-%m-%d %H:%M"), x["tablet"],
             x["a"]["age"], x["a"]["sex"], x["score"], x["glu"], "%.1f" % x["a1c"], x["ref"]]
        rows.append(r)
        if x["event"]["eid"] in resend and x["tablet"].endswith(("3", "7")) and rng.random() < 0.6:
            rows.append(list(r))
    return _csv(path, ["encounter_id", "program_id", "event_id", "screened_at", "tablet_id", "age", "sex", "risk_score",
                       "glucose_mg_dl", "hba1c_pct", "ogtt_referral"], rows)


def write_ogtt(path, T):
    rows = []
    for t in T["tests"]:
        rows.append([t["acc"], t["typed"], t["date"].strftime("%m/%d/%Y"), t["psc"], t["fpg"],
                     "" if t["two"] is None else t["two"], "COMPLETE" if t["complete"] else "INCOMPLETE"])
    return _csv(path, ["ACCESSION", "PROGRAM_ID", "COLLECTED", "PSC", "GLU_FASTING_MGDL",
                       "GLU_120MIN_MGDL", "STATUS"], rows, delimiter="|")


# ---------------------------------------------------------------------------------------------- candidate rules
def write_rules(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidates 2027"
    _head(ws, ["Rule ID", "Rule", "Tablet logic", "Measurements the tablet takes for it"], [10, 52, 50, 46])
    need = {"A60": "HbA1c", "A57": "HbA1c", "R26A60": "Questionnaire, fingerstick glucose, HbA1c",
            "A57G120": "HbA1c, fingerstick glucose", "A55": "HbA1c", "SC5": "Questionnaire"}
    for rid, name, logic in K.RULES:
        ws.append([rid, name, logic, need[rid]])
    ws.append([])
    ws.append([None, "SCORE is the questionnaire's risk score (0 to 10); GLUCOSE the fingerstick random glucose in mg/dL; "
                   "HBA1C the point-of-care result in per cent, as the tablet records them at the screening."])
    ws.cell(ws.max_row, 2).alignment = Alignment(wrap_text=True)
    ws.row_dimensions[ws.max_row].height = 30
    wb.save(path)
    return len(K.RULES)


# ---------------------------------------------------------------------------------------------- close-out
def closeout_counts(T):
    """The season's counts as the close-out reports them: by the month and service area of each adult's first
    screening."""
    tested = {t["x"]["enc"]: t for t in T["tests"]}
    by = {}
    for e in T["E"]:
        k = (e["date"].month, e["area"])
        by.setdefault(k, dict(events=0, adults=0, returns=0, rule=0, draws=0, ogtt=0, found=0))["events"] += 1
    for x in T["enc"]:
        k = (x["event"]["date"].month, x["event"]["area"])
        b = by[k]
        if x["visit"] == 2:
            b["returns"] += 1
            continue
        b["adults"] += 1
        b["rule"] += x["ref"] == "RULE"
        b["draws"] += x["ref"] == "EVAL"
        t = tested.get(x["enc"])
        if t and t["complete"]:
            b["ogtt"] += 1
            b["found"] += int(t["fpg"] >= 126 or (t["two"] is not None and t["two"] >= 200))
    return by


def write_closeout(path, T):
    by = closeout_counts(T)
    cols = ["events", "adults", "returns", "rule", "draws", "ogtt", "found"]
    head = ["Events", "Adults screened", "Return visits", "Referred by the 2026 rule", "Evaluation draws",
            "OGTTs completed", "New diabetes found"]
    wb = Workbook()
    ws = wb.active
    ws.title = "By month"
    ws.append(["%s %s: 2026 season close-out" % (K.ORG, K.PROGRAM)])
    ws.cell(1, 1).font = Font(bold=True, size=12)
    ws.append(["Prepared by %s, data manager, 18 January 2027" % K.PEOPLE["data"]])
    ws.append([])
    _head(ws, ["Month"] + head, [12, 9, 15, 13, 17, 15, 15, 15])
    tot = [0] * len(cols)
    for m in K.MONTHS:
        v = [sum(by[k][c] for k in by if k[0] == m) for c in cols]
        tot = [a + b for a, b in zip(tot, v)]
        ws.append([K.MONTH_NAME[m]] + v)
    ws.append(["Season"] + tot)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws2 = wb.create_sheet("By service area")
    _head(ws2, ["Service area"] + head, [16, 9, 15, 13, 17, 15, 15, 15])
    for a in K.AREAS:
        ws2.append([a] + [sum(by[k][c] for k in by if k[1] == a) for c in cols])
    ws2.append(["Season"] + tot)
    for c in ws2[ws2.max_row]:
        c.font = Font(bold=True)
    ws3 = wb.create_sheet("Notes")
    for line in ["Adults screened counts each adult once, in the month and service area of the first screening of "
                 "the season; return visits are later screenings of the same adults.",
                 "Referrals, draws, OGTTs and new diabetes are counted in the month of the adult's first screening.",
                 "OGTTs completed counts complete tests (both samples taken) received from %s by 15 January 2027."
                 % K.LAB,
                 "New diabetes found: a complete OGTT with a fasting glucose of 126 mg/dL or more or a two-hour glucose "
                 "of 200 mg/dL or more."]:
        ws3.append([line])
    ws3.column_dimensions["A"].width = 120
    wb.save(path)
    return len(K.MONTHS)


# ---------------------------------------------------------------------------------------------- events
def write_events(path, T, use):
    rows = []
    for e in T["E"]:
        end = dt.time(e["start_h"] + e["hours"])
        rows.append([e["eid"], e["date"].isoformat(), e["venue"], e["vtype"], e["area"], "%02d:00" % e["start_h"],
                     end.strftime("%H:%M"), e["tablets"], " ".join(use[e["eid"]])])
    return _csv(path, ["event_id", "event_date", "venue", "venue_type", "service_area", "start_time", "end_time",
                       "tablets", "analyzers"], rows)


# ---------------------------------------------------------------------------------------------- analysers
def write_runlog(path, log):
    """Biomedical Engineering's export of its run database: one row per run it holds, in run-time order."""
    rows = [[r["serial"], r["stamp"].strftime("%Y-%m-%dT%H:%M:%S"), r["type"], r["sample"], r["result"], r["error"],
             r["lot"]] for r in log]
    return _csv(path, ["serial_number", "run_started", "run_type", "sample_id", "result_pct", "error_code",
                       "cartridge_lot"], rows)


def write_assets(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Analyzers"
    _head(ws, ["Asset tag", "Serial number", "Model", "In service from", "In service to", "Note"], [10, 15, 12, 15, 15, 58])
    note = {}
    for tag, (d, new) in AW.REPLACE.items():
        note[(tag, new)] = "Replacement unit; the unit before it went back to the maker under warranty"
    for tag, a, b in AW.LOANS:
        note[(tag, AW.LOANER)] = "Maker's loaner while %s's unit was at the service centre" % tag
    for r in sorted(rows, key=lambda r: (r["tag"], r["frm"])):
        ws.append([r["tag"], r["serial"], "GQ-200", r["frm"], r["to"], note.get((r["tag"], r["serial"]))])
        ws.cell(ws.max_row, 4).number_format = "yyyy-mm-dd"
        ws.cell(ws.max_row, 5).number_format = "yyyy-mm-dd"
    wb.save(path)
    return len(rows)


# ---------------------------------------------------------------------------------------------- invoices
def write_invoices(path, lines):
    """The laboratory's account export: each invoice's lines and each credit memo's lines; a credit memo line carries
    the memo's amount."""
    rows = []
    for l in lines:
        if l["doc"] == "Invoice":
            desc, amount, applies = "Glucose tolerance test, 75 g, fasting and 2 hour", l["amount"], ""
        else:
            desc = "Credit: billed twice" if l["kind"] == "credit_dup" else "Credit: specimen not tested"
            amount = l["memo_total"]
            applies = "%s line %d" % (l["applies"]["number"], l["applies"]["line"])
        rows.append([l["number"], l["date"].isoformat(), l["doc"], l["line"], l["acc"], l["service"].isoformat(),
                     K.OGTT_CODE, desc, "%.2f" % amount, applies])
    return _csv(path, ["document_number", "document_date", "document_type", "line", "accession", "service_date",
                       "test_code", "description", "amount_usd", "applies_to"], rows)


# ---------------------------------------------------------------------------------------------- distractors
def write_feedback(path, T):
    rng = rng_for(9)
    rows = []
    k = 0
    for e in T["E"]:
        n = int(rng.binomial(e["n"], 0.14))
        for j in range(n):
            k += 1
            wait = int(max(2, rng.gamma(2.0, 6.0 + (4 if e["vtype"] == "county fair" else 0))))
            sat = int(np.clip(round(5 - wait / 18 + rng.normal(0, 0.7)), 1, 5))
            rows.append(["FB26-%05d" % k, e["eid"], (e["date"] + dt.timedelta(days=int(rng.integers(0, 4)))).isoformat(),
                         wait, sat, "Y" if sat >= 4 or rng.random() < .2 else "N",
                         ["Flyer", "Church announcement", "Employer", "Radio", "Social media", "Friend or family",
                          "Walked by"][int(rng.choice(7, p=[.14, .22, .16, .06, .14, .18, .10]))]])
    return _csv(path, ["response_id", "event_id", "submitted_on", "wait_minutes", "satisfaction_1_to_5",
                       "would_recommend", "heard_about"], rows)


def write_profile(path):
    rng = rng_for(10)
    wb = Workbook()
    ws = wb.active
    ws.title = "Adults 2024"
    ws.append(["%s region community health assessment, 2024: adults 18 and over" % K.REGION])
    ws.cell(1, 1).font = Font(bold=True, size=12)
    ws.append(["Regional public health office, published September 2025. Survey-based estimates; percentages of adults."])
    ws.append([])
    _head(ws, ["Service area", "Adults 18+", "Diagnosed diabetes %", "Obesity %", "No leisure-time activity %",
               "Current smoking %", "Last checkup within a year %"], [16, 12, 12, 11, 14, 12, 14])
    for a, w in zip(K.AREAS, K.AREA_W):
        ws.append([a, int(round(410000 * w + rng.normal(0, 2500), -1)), round(float(rng.uniform(9.4, 13.8)), 1),
                   round(float(rng.uniform(31.0, 39.5)), 1), round(float(rng.uniform(22.0, 31.0)), 1),
                   round(float(rng.uniform(11.0, 17.5)), 1), round(float(rng.uniform(70.0, 79.0)), 1)])
    wb.save(path)
    return len(K.AREAS)
