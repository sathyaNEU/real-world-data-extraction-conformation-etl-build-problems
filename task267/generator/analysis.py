"""Every figure from the shipped pack alone: the adults screened (first screenings), their OGTTs, each candidate
rule's referral share and its sensitivity under each rung's estimator, the call, and the two asks with their stops."""
import csv
import datetime as dt
import os
import re
from collections import defaultdict

import numpy as np
from openpyxl import load_workbook

import common as K


def _csv(path, delimiter=","):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh, delimiter=delimiter))


def _sheet(path, name, header_row=1):
    ws = load_workbook(path, read_only=True, data_only=True)[name]
    it = ws.iter_rows(values_only=True)
    for _ in range(header_row - 1):
        next(it)
    head = next(it)
    return [dict(zip(head, r)) for r in it if any(v is not None for v in r)]


def pid_key(s):
    d = re.sub(r"\D", "", s or "")
    return "SF-%07d" % int(d) if d else None


class Pack:
    def __init__(self, target, keep_duplicates=False, every_encounter=False):
        P = lambda k: os.path.join(target, K.FILES[k])
        rows = _csv(P("encounters"))
        if not keep_duplicates:
            seen, out = set(), []
            for r in rows:
                if r["encounter_id"] in seen:
                    continue
                seen.add(r["encounter_id"])
                out.append(r)
            rows = out
        rows.sort(key=lambda r: (r["screened_at"], r["encounter_id"]))
        if every_encounter:
            self.adults = rows
        else:
            first = {}
            for r in rows:
                first.setdefault(r["program_id"], r)
            self.adults = sorted(first.values(), key=lambda r: r["encounter_id"])
        self.encounters = rows
        tests = _csv(P("ogtt"), delimiter="|")
        self.tests = tests
        res = {}
        for t in tests:
            if t["STATUS"] != "COMPLETE":
                continue
            k = pid_key(t["PROGRAM_ID"])
            dm = int(t["GLU_FASTING_MGDL"]) >= 126 or int(t["GLU_120MIN_MGDL"]) >= 200
            assert k not in res, k
            res[k] = dm
        self.result = res
        ev = {e["event_id"]: e for e in _csv(P("events"))}
        self.events = ev
        a = self.adults
        self.score = np.array([int(r["risk_score"]) for r in a])
        self.glu = np.array([int(r["glucose_mg_dl"]) for r in a])
        self.a1c = np.array([float(r["hba1c_pct"]) for r in a])
        self.ref = np.array([r["ogtt_referral"] for r in a])
        self.ver = np.array([r["program_id"] in res for r in a])
        self.dm = np.array([bool(res.get(r["program_id"], False)) for r in a])
        self.area = np.array([ev[r["event_id"]]["service_area"] for r in a])
        self.month = np.array([int(r["screened_at"][5:7]) for r in a])


def weights(pk, groups):
    w = np.zeros(len(pk.adults))
    for g in sorted(set(groups)):
        m = groups == g
        v = m & pk.ver
        if v.sum():
            w[v] = m.sum() / v.sum()
    return w


def sens(pk, f, w):
    return 100 * (w * (f & pk.dm & pk.ver)).sum() / (w * (pk.dm & pk.ver)).sum()


def rungs(pk):
    """Each rule's referral share and its sensitivity under every rung, per cent."""
    rule = pk.ref == "RULE"
    W = dict(golden=weights(pk, np.where(rule, "R", "N")),
             area=weights(pk, np.char.add(np.where(rule, "R", "N"), pk.area.astype(str))),
             month=weights(pk, np.char.add(np.where(rule, "R", "N"), pk.month.astype(str))),
             nominal=np.where(pk.ver, np.where(rule, 1.0, 1 / K.EVAL_FRACTION), 0.0),
             cc=pk.ver.astype(float))
    ev = (pk.ref == "EVAL") & pk.ver
    out = {}
    for r in K.RULE_IDS:
        f = K.flags(r, pk.score, pk.glu, pk.a1c)
        o = dict(share=100 * f.mean())
        for k, w in W.items():
            o[k] = sens(pk, f, w)
        o["untested_free"] = 100 * (f & pk.dm & pk.ver).sum() / (pk.dm & pk.ver).sum()
        o["evalonly"] = 100 * (f & pk.dm & ev).sum() / (pk.dm & ev).sum()
        pT = f.mean()
        p1 = (pk.dm & pk.ver & f).sum() / (pk.ver & f).sum()
        p0 = (pk.dm & pk.ver & ~f).sum() / (pk.ver & ~f).sum()
        o["own_rates"] = 100 * pT * p1 / (pT * p1 + (1 - pT) * p0)
        # each rule's flags crossed with the 2026 group, each cell's diabetes rate among its tested adults
        num = den = 0.0
        for g in (True, False):
            for fv in (True, False):
                m = (rule == g) & (f == fv)
                v = m & pk.ver
                if v.sum():
                    p = (pk.dm & v).sum() / v.sum()
                    den += m.sum() * p
                    num += m.sum() * p * fv
        o["cells"] = 100 * num / den
        out[r] = o
    return out


def call(R, key="golden"):
    ok = [r for r in K.RULE_IDS if R[r][key] >= K.TARGET]
    return min(ok, key=lambda r: R[r]["share"]) if ok else "2026 rule"


# ---------------------------------------------------------------------------------------------- Ask A
FIELDS = ("serial_number", "run_type", "sample_id", "result_pct", "error_code", "cartridge_lot")


def runlog(target, twins="drop"):
    """The run database's rows with their times parsed. twins: 'drop' removes a row whose every field but the run
    time repeats a row exactly four or five hours earlier on the same unit (the same run loaded once in local time
    and once in UTC); 'keep' leaves them; 'sameday' counts rows identical but for the run time once a day."""
    rows = _csv(os.path.join(target, K.FILES["runlog"]))
    for r in rows:
        r["t"] = dt.datetime.fromisoformat(r["run_started"])
    if twins == "keep":
        return rows
    if twins == "sameday":
        seen, out = set(), []
        for r in rows:
            k = tuple(r[f] for f in FIELDS) + (r["t"].date(),)
            if k in seen:
                continue
            seen.add(k)
            out.append(r)
        return out
    held = {(tuple(r[f] for f in FIELDS), r["t"]) for r in rows}
    return [r for r in rows if not any((tuple(r[f] for f in FIELDS), r["t"] - dt.timedelta(hours=h)) in held
                                       for h in (4, 5))]


def register(target):
    reg = _sheet(os.path.join(target, K.FILES["assets"]), "Analyzers")
    for r in reg:
        r["frm"] = r["In service from"].date() if hasattr(r["In service from"], "date") else r["In service from"]
        to = r["In service to"]
        r["to"] = to.date() if hasattr(to, "date") else to
    return reg


def spent_codes(target):
    from pypdf import PdfReader
    txt = "".join(p.extract_text() for p in PdfReader(os.path.join(target, K.FILES["errors"])).pages)
    spent = {m.group(1) for m in re.finditer(r"(E-\d\d)\s+[^\n]+\s+During the measurement", txt)}
    assert spent, "no measurement-phase codes read from the reference"
    return spent


def cartridges(target, twins="drop", errors="rule", window=True, by="tag"):
    """Cartridges each analyzer used over the season. errors: 'rule' (a run ended at the cartridge check uses none),
    'all' (every error row uses one), 'results' (only rows with a result). by: 'tag' (the unit behind each asset tag
    that day) or 'serial'."""
    log = runlog(target, twins)
    reg = register(target)
    spent = spent_codes(target)
    out = defaultdict(int)
    for r in log:
        d = r["t"].date()
        if window and not (K.SEASON_START <= d <= K.SEASON_END):
            continue
        e = r["error_code"]
        if errors == "results" and e:
            continue
        if errors == "rule" and e and e not in spent:
            continue
        if by == "serial":
            out[r["serial_number"]] += 1
            continue
        tags = [a["Asset tag"] for a in reg if a["Serial number"] == r["serial_number"] and a["frm"] <= d and
                (a["to"] is None or d <= a["to"])]
        assert len(tags) == 1, (r["serial_number"], d, tags)
        out[tags[0]] += 1
    return dict(out)


def ask_a(target):
    return cartridges(target)


def ask_a_stops(target):
    return {"every row on its tag, nothing removed": cartridges(target, twins="keep", errors="all", window=False),
            "error rule and season applied, every row kept": cartridges(target, twins="keep"),
            "result rows only, every row kept": cartridges(target, twins="keep", errors="results"),
            "rows alike but for the run time counted once a day": cartridges(target, twins="sameday"),
            "twins removed, every error counted": cartridges(target, errors="all"),
            "twins removed, the training day kept": cartridges(target, window=False),
            "answer": cartridges(target)}


# ---------------------------------------------------------------------------------------------- Ask B
def ask_b(target, month_of="test", credits="applied", program="extract", dedupe=False):
    """For each month of the season, in cents: what the laboratory billed net of its credits for the tests done that
    month, and the part for tests of screened adults (accessions in the program's OGTT results). month_of: 'test' (a
    charge by its service date, a credit by the service date of the invoice line it applies to), 'row' (every row by
    its own service_date column, a credit memo line falling on the memo's date), 'document' (the month before the
    document's date). credits: 'applied' or 'none'."""
    P = lambda k: os.path.join(target, K.FILES[k])
    lines = _csv(P("invoices"))
    prog = {t["ACCESSION"] for t in _csv(P("ogtt"), delimiter="|")}
    by_line = {(l["document_number"], l["line"]): l for l in lines}
    out = {m: [0, 0] for m in K.MONTHS}
    seen = set()
    for l in lines:
        credit = l["document_type"] != "Invoice"
        if credit and credits == "none":
            continue
        if dedupe and not credit:
            if l["accession"] in seen:
                continue
            seen.add(l["accession"])
        v = int(round(float(l["amount_usd"]) * 100))
        if month_of == "test":
            src = by_line[tuple(l["applies_to"].rsplit(" line ", 1))] if credit else l
            d = dt.date.fromisoformat(src["service_date"])
        elif month_of == "row":
            d = dt.date.fromisoformat(l["service_date"])
        else:
            d = dt.date.fromisoformat(l["document_date"]).replace(day=1) - dt.timedelta(days=1)
        if d.year == 2026 and d.month in out:
            out[d.month][0] += v
            out[d.month][1] += v if (program == "all" or l["accession"] in prog) else 0
    return out


def ask_b_stops(target):
    return {"charges by the month billed": ask_b(target, month_of="document", credits="none"),
            "charges by the month done, credits left out": ask_b(target, credits="none"),
            "every row by its own date column, credits on the memo's date": ask_b(target, month_of="row"),
            "every row by the month billed": ask_b(target, month_of="document"),
            "one charge per accession, credits applied": ask_b(target, dedupe=True),
            "every billed test counted as the program's": ask_b(target, program="all"),
            "answer": ask_b(target)}
