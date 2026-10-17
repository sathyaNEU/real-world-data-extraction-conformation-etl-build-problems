"""The goldens: referral_rule_review_2027.xlsx (each rule's two shares, the analyzers' cartridges, the laboratory's
billing by month) and referral_rule_2027_memo.docx (the analyst's memo to the program director with the chart), plus
generator/figures.json and metadata.json."""
import json
import os
import shutil
import tempfile

import analysis as A
import common as K
import postprocess as PP

HERE = os.path.dirname(os.path.abspath(__file__))
CHART = "referral_rules_2027.png"
LABP = K.LAB + "'"
SHORT = {"A60": "HbA1c 6.0", "A57": "HbA1c 5.7", "R26A60": "2026 rule or HbA1c 6.0", "A57G120": "HbA1c 5.7 or glucose 120",
         "A55": "HbA1c 5.5", "SC5": "Risk score 5"}


def r1(x):
    return int(x * 10 + 0.5) / 10


def r0(x):
    """A sensitivity, an estimate, to the whole percent."""
    return int(x + 0.5)


def results(target):
    pk = A.Pack(target)
    R = A.rungs(pk)
    rule = pk.ref == "RULE"
    w = A.weights(pk, __import__("numpy").where(rule, "R", "N"))
    est_dm = (w * (pk.dm & pk.ver)).sum()
    est_dm_not_rule = (w * (pk.dm & pk.ver & ~rule)).sum()
    return dict(pk=pk, R=R, chosen=A.call(R), default=A.call(R, "cc"), evalonly=A.call(R, "evalonly"),
                adults=len(pk.adults), tested=int(pk.ver.sum()), rule_n=int(rule.sum()),
                draws=int((pk.ref == "EVAL").sum()), not_rule_share=100 * est_dm_not_rule / est_dm,
                ask_a=A.ask_a(target), ask_b=A.ask_b(target))


# ---------------------------------------------------------------------------------------------- the workbook
def _note(ws, text, ncols, height):
    from openpyxl.styles import Alignment
    ws.append([text])
    r = ws.max_row
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
    ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = height


def _print(ws, title_row, landscape=False):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "%d:%d" % (title_row, title_row)


def workbook(path, o):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    R = o["R"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Referral rules"
    head_fill = PatternFill("solid", fgColor="2F5D50")
    thin = Side(style="thin", color="B7BEC4")
    ws.append(["HbA1c of 5.7% or higher, or random glucose of 120 mg/dL or higher, is the 2027 rule: the fewest "
               "referrals among the rules that reach 85%"])
    ws.cell(1, 1).font = Font(bold=True, size=12)
    ws.merge_cells("A1:F1")
    ws.cell(1, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[1].height = 32
    ws.append(["Protocol 2026-27, section 6: a rule qualifies at 85% or more of adults with undiagnosed diabetes; the "
               "qualifying rule that refers the fewest adults is used."])
    ws.cell(2, 1).font = Font(italic=True, color="5B646B")
    ws.merge_cells("A2:F2")
    ws.append([])
    head = ["Rule ID", "Rule", "Adults with undiagnosed diabetes it would have referred", "Adults screened it would "
            "have referred", "Meets the 85% target", "2027 rule"]
    ws.append(head)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[4].height = 44
    for r in sorted(K.RULE_IDS, key=lambda r: R[r]["share"]):
        ws.append([r, K.RULE_NAME[r], r0(R[r]["golden"]) / 100, r1(R[r]["share"]) / 100,
                   "Yes" if R[r]["golden"] >= K.TARGET else "No", "Chosen" if r == o["chosen"] else None])
        row = ws.max_row
        ws.cell(row, 3).number_format = "0%"
        ws.cell(row, 4).number_format = "0.0%"
        for c in (5, 6):
            ws.cell(row, c).alignment = Alignment(horizontal="center")
        if r == o["chosen"]:
            for c in ws[row]:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="E3EFEA")
        for c in ws[row]:
            c.border = Border(bottom=thin)
    for col, wdt in zip("ABCDEF", (10, 58, 22, 20, 12, 11)):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "A5"
    ws.append([])
    _note(ws, "Shares among the %s adults screened in 2026 (first screening of the season). The share of adults "
              "with undiagnosed diabetes counts every adult screened: each adult with a complete OGTT stands for the "
              "adults of the same 2026 group (referred by the 2026 rule, or not referred), in proportion to the "
              "group's complete tests." % "{:,}".format(o["adults"]), 6, 44)
    _print(ws, 4, landscape=True)

    ws2 = wb.create_sheet("Analyzer cartridges")
    ws2.append(["Asset tag", "HbA1c cartridges used, 2026 season"])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True)
    for t in K.ASSETS:
        ws2.append([t, o["ask_a"][t]])
        ws2.cell(ws2.max_row, 2).number_format = "#,##0"
    ws2.append(["All twelve", sum(o["ask_a"].values())])
    ws2.cell(ws2.max_row, 2).number_format = "#,##0"
    for c in ws2[ws2.max_row]:
        c.font = Font(bold=True)
    ws2.column_dimensions["A"].width = 12
    ws2.column_dimensions["B"].width = 22
    ws2.append([])
    for line in ["Runs on the season's days, 2 March to 20 November 2026, each on the asset tag whose unit it was that day "
                 "(asset register). A run stopped at the cartridge check takes no cartridge; controls take one.",
                 "The run database holds some runs twice, once in local time and once in UTC (four hours later in "
                 "summer, five before 8 March), with every other field the same; each is counted once."]:
        _note(ws2, line, 5, 46)
    ws2.freeze_panes = "A2"
    _print(ws2, 1)

    ws3 = wb.create_sheet("Laboratory billing")
    ws3.append(["Month the tests were done", "Billed, net of credits", "Of which, tests of adults we screened"])
    for c in ws3[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True)
    money = "$#,##0.00"
    for m in K.MONTHS:
        ws3.append(["%s 2026" % K.MONTH_NAME[m], o["ask_b"][m][0] / 100, o["ask_b"][m][1] / 100])
        for c in (2, 3):
            ws3.cell(ws3.max_row, c).number_format = money
    ws3.append(["March to November", sum(v[0] for v in o["ask_b"].values()) / 100,
                sum(v[1] for v in o["ask_b"].values()) / 100])
    for c in ws3[ws3.max_row]:
        c.font = Font(bold=True)
    for c in (2, 3):
        ws3.cell(ws3.max_row, c).number_format = money
    ws3.column_dimensions["A"].width = 26
    ws3.column_dimensions["B"].width = 20
    ws3.column_dimensions["C"].width = 24
    ws3.append([])
    for line in ["%s account export, invoices and credit memos to 6 January 2027, by the service date of each test." % LABP,
                 "A credit memo line carries the memo's whole amount; each line is taken at the amount of the invoice "
                 "line it applies to.",
                 "Tests of adults we screened are the accessions in the program's OGTT results; the rest were billed to "
                 "the program's account for other patients."]:
        _note(ws3, line, 3, 30)
    ws3.freeze_panes = "A2"
    _print(ws3, 1)

    ws4 = wb.create_sheet("Notes")
    for line in ["Prepared by %s, %s, 22 January 2027." % (K.PEOPLE["analyst"], K.PROGRAM),
                 "Adults screened: each adult once, at the first screening of the 2026 season (%s adults)."
                 % "{:,}".format(o["adults"]),
                 "Undiagnosed diabetes: a complete OGTT with a fasting glucose of 126 mg/dL or more or a two-hour "
                 "glucose of 200 mg/dL or more (protocol, section 5).",
                 "%s adults had a complete OGTT: %s referred by the 2026 rule were offered the test, and %s were "
                 "offered it in the evaluation draw." % ("{:,}".format(o["tested"]), "{:,}".format(o["rule_n"]),
                                                         "{:,}".format(o["draws"])),
                 "Shares of the adults with undiagnosed diabetes are estimates, given to the whole percent; shares of the "
                 "adults screened are counts, given to one decimal place; both are rounded from unrounded figures."]:
        ws4.append([line])
        ws4.cell(ws4.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws4.column_dimensions["A"].width = 110
    _print(ws4, 1)
    wb.save(path)


# ---------------------------------------------------------------------------------------------- the chart
def chart(path, o):
    import matplotlib
    matplotlib.use("Agg")
    matplotlib.rcParams["svg.hashsalt"] = "task267"
    import matplotlib.pyplot as plt
    R = o["R"]
    INK, MUTED, GRID = "#1d2326", "#5b646b", "#e3e6e8"
    BLUE, ORANGE = "#2a78d6", "#eb6834"
    fig, ax = plt.subplots(figsize=(7.2, 4.1), dpi=200)
    for r in K.RULE_IDS:
        x, y0, y1 = R[r]["share"], R[r]["cc"], R[r]["golden"]
        ax.plot([x, x], [y0, y1], color="#b7bec4", lw=1.0, zorder=1)
    xs = [R[r]["share"] for r in K.RULE_IDS]
    ax.scatter(xs, [R[r]["cc"] for r in K.RULE_IDS], s=42, facecolors="white", edgecolors=ORANGE, linewidths=1.6,
               zorder=3, label="Among adults with a complete OGTT")
    ax.scatter(xs, [R[r]["golden"] for r in K.RULE_IDS], s=46, color=BLUE, edgecolors="white", linewidths=1.2, zorder=4,
               label="Among everyone screened")
    c = o["chosen"]
    ax.scatter([R[c]["share"]], [R[c]["golden"]], s=150, facecolors="none", edgecolors=INK, linewidths=1.4, zorder=5)
    ax.axhline(K.TARGET, color=INK, lw=1.0, ls=(0, (4, 3)), zorder=2)
    ax.text(4.6, K.TARGET + 0.9, "Target, 85%", color=INK, fontsize=7.6, ha="left")
    off = {"A60": (0.8, -3.2), "A57": (0.8, -3.6), "R26A60": (-0.8, -4.0), "A57G120": (0.9, -3.4), "A55": (-0.8, 2.6),
           "SC5": (0.8, -1.2)}
    for r in K.RULE_IDS:
        dx, dy = off[r]
        ax.text(R[r]["share"] + dx, R[r]["golden"] + dy, SHORT[r] + (" (chosen)" if r == c else ""), fontsize=7.4,
                color=INK if r == c else MUTED, ha="left" if dx > 0 else "right", fontweight="bold" if r == c else None)
    ax.set_xlim(4, 52)
    ax.set_ylim(30, 100)
    ax.set_xlabel("Adults screened the rule would have referred (%)", fontsize=8, color=MUTED)
    ax.set_ylabel("Adults with undiagnosed diabetes referred (%)", fontsize=8, color=MUTED)
    ax.tick_params(labelsize=7.5, colors=MUTED)
    ax.grid(color=GRID, lw=0.6)
    ax.set_axisbelow(True)
    for s_ in ("top", "right"):
        ax.spines[s_].set_visible(False)
    for s_ in ("left", "bottom"):
        ax.spines[s_].set_color("#b7bec4")
    ax.legend(loc="lower right", fontsize=7.4, frameon=False)
    ax.set_title("Read among the tested, the 2026 rule plus HbA1c 6.0 looks like it qualifies; among everyone "
                 "screened it does not", fontsize=8.6, color=INK, loc="left")
    fig.subplots_adjust(left=.09, right=.98, top=.9, bottom=.13)
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)


# ---------------------------------------------------------------------------------------------- the memo
def memo(path, o, png):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt, RGBColor
    R, c = o["R"], o["chosen"]
    GREY = RGBColor(0x5B, 0x64, 0x6B)
    doc = docx.Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(9.5)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(215.9), Mm(279.4)
    sec.left_margin = sec.right_margin = Mm(19)
    sec.top_margin, sec.bottom_margin = Mm(13), Mm(12)
    hp = sec.header.paragraphs[0]
    hp.text = "%s | %s | 2027 season planning" % (K.ORG, K.PROGRAM)
    hp.runs[0].font.size = Pt(7.5)
    hp.runs[0].font.color.rgb = GREY

    def para(text="", size=None, bold=False, color=None, after=3, before=0):
        p = doc.add_paragraph()
        if text:
            r = p.add_run(text)
            r.bold = bold
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        return p

    for k, v in (("To", "%s, Director, %s" % (K.PEOPLE["director"], K.PROGRAM)),
                 ("Cc", "%s, MD, Medical Director" % K.PEOPLE["medical"]),
                 ("From", "%s, program analyst" % K.PEOPLE["analyst"]),
                 ("Date", "22 January 2027"),
                 ("Re", "Referral rule for the 2027 season (protocol 2026-27, section 6)")):
        p = para(after=0)
        r = p.add_run("%s:\t" % k)
        r.bold = True
        p.add_run(v)
    para("HbA1c of 5.7% or higher, or a random glucose of 120 mg/dL or higher, is the rule for 2027", 13, bold=True,
         before=8, after=4)
    p = para(after=4)
    p.add_run("Load rule %s into the tablets: refer every adult with an HbA1c of 5.7%% or higher or a random glucose "
              "of 120 mg/dL or higher. " % c).bold = True
    p.add_run("Among everyone we screened in 2026 it would have referred %d%% of the adults with undiagnosed "
              "diabetes, above the protocol's 85%%, while referring %.1f%% of the adults screened. No rule that refers "
              "fewer adults reaches 85%%: an HbA1c of 5.7%% or higher alone finds %d%%, and the 2026 rule plus an "
              "HbA1c of 6.0%% or higher %d%%." % (r0(R[c]["golden"]), r1(R[c]["share"]), r0(R["A57"]["golden"]),
                                                  r0(R["R26A60"]["golden"])))
    tb = doc.add_table(rows=1, cols=4)
    tb.style = "Light List"
    for i, t in enumerate(["Candidate rule", "Adults screened referred", "Undiagnosed diabetes referred",
                           "Meets 85%"]):
        tb.rows[0].cells[i].text = t
    for r in sorted(K.RULE_IDS, key=lambda r: R[r]["share"]):
        cells = tb.add_row().cells
        cells[0].text = "%s (%s)" % (K.RULE_NAME[r], r)
        cells[1].text = "%.1f%%" % r1(R[r]["share"])
        cells[2].text = "%d%%" % r0(R[r]["golden"])
        cells[3].text = "Yes" if R[r]["golden"] >= K.TARGET else "No"
        for i in (1, 2):
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if r == c:
            for cell in cells:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    tb.autofit = False
    for i, w in enumerate((98, 27, 33, 19)):
        tb.columns[i].width = Mm(w)
        for cell in tb.columns[i].cells:
            cell.width = Mm(w)
    para("Source: screening encounters and the OGTT results from %s, 2026 season; %s adults screened, %s with a complete "
         "OGTT." % (K.LAB, "{:,}".format(o["adults"]), "{:,}".format(o["tested"])), 7.5, color=GREY, before=2, after=6)
    para("Most of the season's OGTTs went to adults the 2026 rule referred; the evaluation draw offered the test to one "
         "in five of everyone else. Read among the adults who were tested, the 2026 rule plus an HbA1c of 6.0%% or "
         "higher finds %d%% of the diabetes, because nearly every tested case was one the 2026 rule had already "
         "flagged. Once each tested adult stands for the adults of their own group, %.0f%% of the adults with "
         "undiagnosed diabetes were among those the 2026 rule did not refer, milder cases that an HbA1c of 6.0%% "
         "misses, and that rule falls to %d%%." % (r0(R["R26A60"]["cc"]), o["not_rule_share"],
                                                   r0(R["R26A60"]["golden"])), after=4)
    doc.add_picture(png, width=Mm(146))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para("The season's records are in the review workbook: the HbA1c cartridges each analyzer used, %s in all with "
         "the replaced units and the loaner counted on the tags they served, and %s billing by month, $%s net for "
         "the season's OGTTs, $%s of it for adults we screened." % (
             "{:,}".format(sum(o["ask_a"].values())), LABP, "{:,.2f}".format(sum(v[0] for v in o["ask_b"].values()) / 100),
             "{:,.2f}".format(sum(v[1] for v in o["ask_b"].values()) / 100)), before=4, after=3)
    para("Each adult counts once, at the first screening of the season. Shares of the adults with undiagnosed diabetes "
         "are estimates, given to the whole percent; the target is compared unrounded.", 7.5, color=GREY, after=0)
    doc.save(path)


# ---------------------------------------------------------------------------------------------- figures.json
def figures(o):
    R = o["R"]
    return dict(
        call=dict(rule=o["chosen"], rule_name=K.RULE_NAME[o["chosen"]], target=K.TARGET,
                  sensitivity={K.RULE_NAME[r]: r0(R[r]["golden"]) for r in K.RULE_IDS},
                  referral_share={K.RULE_NAME[r]: r1(R[r]["share"]) for r in K.RULE_IDS},
                  qualifies=[K.RULE_NAME[r] for r in K.RULE_IDS if R[r]["golden"] >= K.TARGET]),
        rungs={k: {K.RULE_NAME[r]: r0(R[r][k]) for r in K.RULE_IDS} for k in
               ("cc", "untested_free", "evalonly", "own_rates", "cells", "area", "month", "nominal")},
        rung_picks={k: A.call(R, k) for k in ("cc", "untested_free", "evalonly", "own_rates", "cells", "area", "month",
                                              "nominal", "golden")},
        counts=dict(adults=o["adults"], tested=o["tested"], referred_2026=o["rule_n"], draws=o["draws"],
                    not_rule_share=int(o["not_rule_share"] + 0.5)),
        ask_a={t: o["ask_a"][t] for t in K.ASSETS}, ask_a_total=sum(o["ask_a"].values()),
        ask_b={K.MONTH_NAME[m]: {"billed": o["ask_b"][m][0] / 100, "screened": o["ask_b"][m][1] / 100}
               for m in K.MONTHS},
        ask_b_totals=dict(billed=sum(v[0] for v in o["ask_b"].values()) / 100,
                          screened=sum(v[1] for v in o["ask_b"].values()) / 100),
        protocol=dict(sections=["2.2", "4.1", "4.2", "5.1", "6.2", "6.3"], fasting=126, two_hour=200, score=6,
                      glucose=140, fifth=5, season_year=2026, next_season=2027, decimals=1),
        rule_thresholds=dict(hba1c=[6.0, 5.7, 5.5], glucose=120, score=5))


def write(target, golden, T=None):
    os.makedirs(golden, exist_ok=True)
    o = results(target)
    xp = os.path.join(golden, K.DELIVERABLES[0])
    workbook(xp, o)
    tmp = tempfile.mkdtemp()
    png = os.path.join(tmp, CHART)
    chart(png, o)
    dp = os.path.join(golden, K.DELIVERABLES[1])
    memo(dp, o, png)
    shutil.rmtree(tmp)
    PP.COMPANY = K.ORG
    PP.normalize_ooxml(xp, K.PEOPLE["analyst"], K.PEOPLE["analyst"], "2027-01-19T14:22:05Z", "2027-01-22T19:48:31Z",
                       excel=True)
    PP.normalize_ooxml(dp, K.PEOPLE["analyst"], K.PEOPLE["analyst"], "2027-01-20T15:03:44Z", "2027-01-22T20:11:09Z",
                       title="Referral rule for the 2027 season")
    json.dump(figures(o), open(os.path.join(HERE, "figures.json"), "w"), indent=1, sort_keys=True)
    meta = {"distractor_files": ["target/" + K.FILES[k] for k in K.DISTRACTORS],
            "deliverables": K.DELIVERABLES, "stated_inputs": []}
    json.dump(meta, open(os.path.join(os.path.dirname(HERE), "metadata.json"), "w"), indent=2)
    return o
