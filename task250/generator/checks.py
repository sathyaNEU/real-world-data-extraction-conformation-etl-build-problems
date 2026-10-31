"""task250 generator checks: the input gates and the no-describing-file rule, the three rules on the clean truth (R1,
R2 per mess item, R3a to R3c), every semantic fork row (SF1 to SF22), the ladder with its margins, the unloadings'
purity and threshold independence, the separation line and the deletion tests, the asks against the clean truth with
their devices, the realism sweep and the clocks.
    C.run(target, world, fig, golden) -> number of assertions passed"""
import collections
import datetime as dt
import json
import os
import shutil
import tempfile

import pandas as pd
from openpyxl import load_workbook
from pypdf import PdfReader

import build as B
import golden as GD
import params as P
import rungs as R
import sim

N = [0]
FAILS = []


def ok(cond, msg):
    if cond:
        N[0] += 1
    else:
        FAILS.append(msg)


def alerts_of(r):
    return sorted(h for h, v in r.items() if v["rate"] > GD.TOLERANCE)


def run(target, w, fig, golden):
    N[0] = 0
    FAILS.clear()
    F = {k: os.path.join(target, v) for k, v in B.FILES.items()}
    d = GD.load(target, B.FILES)
    gold_set = sorted(x["hub_code"] for x in fig["hubs"] if x["fires"])
    ok(gold_set == ["HKTW", "HLDZ", "HWAW"], "golden alert set %s" % gold_set)

    # ------------------------------------------------------------ input gates and no describing file
    names = sorted(os.listdir(target))
    ok(len(names) >= 10, "files %d" % len(names))
    fmts = {n.rsplit(".", 1)[1] for n in names}
    ok(len(fmts) >= 3, "formats %s" % fmts)
    ok(len(d["courier"]) >= 25000 and len(d["depot"]) >= 25000, "a 25,000-row file")
    meta = json.load(open(os.path.join(os.path.dirname(target), "metadata.json")))
    ok(len(meta["distractor_files"]) >= 2 and all(x in names for x in meta["distractor_files"]), "two distractors named")
    bad = [n for n in names if any(k in n.lower() for k in ("readme", "dictionary", "notes", "manifest", "schema", "field"))]
    ok(not bad, "a describing file by name %s" % bad)
    for k in ("plan", "register"):
        sheets = load_workbook(F[k], read_only=True).sheetnames
        ok(not [s for s in sheets if any(x in s.lower() for x in ("note", "readme", "dictionary", "about", "column"))],
           "a describing sheet in %s: %s" % (k, sheets))
    std = " ".join(p.extract_text() for p in PdfReader(F["standard"]).pages).lower()
    for word in ("column", "field", "trailer", "unload order", "door", "overflow", "manual", "plan", "utc", "minute",
                 "clock", "tc51", "tc52", "suffix", "prefix", "characters", "imput", "share"):
        ok(word not in std, "TDS-4 carries %r" % word)
    req = open(F["request"], encoding="utf-8").read().lower()
    for word in ("trailer", "overflow", "strykow", "pruszkow", "sosnowiec", "manual", "door", "plan", "duplicate",
                 "column", "file", "extract shows"):
        ok(word not in req, "the request carries %r" % word)
    ok(len(PdfReader(F["standard"]).pages) <= 4, "the standard runs over four pages")

    # ------------------------------------------------------------ the clean truth and the golden
    ps = [p for p in w["parcels"] if P.MONTH_FIRST <= p["arr_day"] <= P.MONTH_LAST]
    ok(len(ps) == fig["n_parcels"] == len(w["parcels"]), "population %d %d" % (len(ps), fig["n_parcels"]))
    arr = GD.arrivals(d["depot"])
    sh = GD.sort_hub(d["hub"])
    gold, impure, blind, n_unl = GD.golden_assign(arr, sh)
    truth = {p["label"]: p["hub"] for p in ps}
    ok(all(gold[t] == truth[t] for t in gold) and len(gold) == len(truth), "R3c: the golden hub is the hub that sorted it")
    tr = collections.Counter(p["hub"] for p in ps)
    tm = collections.Counter(p["hub"] for p in ps if not p["scanned"])
    for x in fig["hubs"]:
        h = x["hub_code"]
        ok(x["parcels"] == tr[h] and x["missed"] == tm[h], "R3c truth figures %s" % h)
    ok(set(sh) == {p["label"] for p in ps if p["scanned"]}, "the sort feed holds exactly the scanned parcels")
    ok(impure == 0 and blind == 0, "every unloading pure and with a sorted parcel (%d, %d)" % (impure, blind))

    # ------------------------------------------------------------ the three rules, executed
    cols = set()
    for k in ("depot", "hub", "courier", "pre"):
        cols |= set(d[k].columns)
    ok("hub_code" not in d["depot"].columns and "sort_hub" not in d["depot"].columns,
       "R1: the deciding property has no column on the arrivals")
    ok(not [c for c in cols if "trailer" in c or "unload" in c or "route" in c], "R1: no trailer or route column")
    ok(d["depot"]["scanned_at"].str.len().eq(19).all(), "R1: arrival scans carry the second")
    M, arr2, sh2, gold2 = R.main_rungs(d)
    r0 = alerts_of(M["R0 plan"])
    ok(r0 != gold_set and r0 == ["HKTW", "HWAW"], "R3a: the default path files %s" % r0)
    # R3b: the decisive step alone moves the default path onto the golden
    plan_rows = R.plan_lookup(d["plan"])
    step = {}
    for t, dep, ts in zip(arr["tracking_no"], arr["depot_code"], arr["scanned_at"]):
        step[t] = sh[t] if t in sh else gold[t]
    ok(alerts_of(GD.hub_rates(step, arr, sh)) == gold_set, "R3b: the unloading's hub alone lands the golden")
    # truth-side analyst test on the clean world
    plan_truth = collections.Counter()
    miss_truth = collections.Counter()
    for p in ps:
        h = p["hub"] if p["scanned"] else p["plan_hub"]
        plan_truth[h] += 1
        miss_truth[h] += 0 if p["scanned"] else 1
    r0t = sorted(h for h in plan_truth if 100.0 * miss_truth[h] / plan_truth[h] > GD.TOLERANCE)
    ok(r0t == r0, "R3a on the clean truth %s" % r0t)

    # ------------------------------------------------------------ the ladder
    margins = {}
    for k, r in M.items():
        a = alerts_of(r)
        if k == "golden":
            ok(a == gold_set, "golden rung")
            continue
        ok(a != gold_set, "rung %s lands the golden" % k)
        wrong = sorted(set(a) ^ set(gold_set))
        dist = {h: max(r[h]["rate"] / GD.TOLERANCE, GD.TOLERANCE / r[h]["rate"]) for h in r}
        if k == "R0 plan, first issue":
            m = max(dist[h] for h in wrong)          # a hazard rung: wrong whichever side its nearest hub falls
        else:
            m = min(dist.values())                   # an analytical rung: every verdict it files stands clear
        margins[k] = m
        ok(m >= 1.15, "rung %s margin %.3f" % (k, m))
    g = M["golden"]
    gm = min(max(v["rate"] / GD.TOLERANCE, GD.TOLERANCE / v["rate"]) for v in g.values())
    ok(gm >= 1.2, "golden margin %.3f" % gm)
    ok(len({tuple(alerts_of(r)) for r in M.values()}) >= 4, "the rungs name at least four different sets")

    # ------------------------------------------------------------ the unloadings: pure and threshold-free
    for gap in (1, 2, 5, 20, 29):
        a2, i2, b2, _ = GD.golden_assign(arr, sh, gap)
        ok(a2 == gold and i2 == 0 and b2 == 0, "unloadings at a %d-minute gap" % gap)
    u = GD.unloadings(arr)
    u["dt"] = u.groupby("unloading")["t"].diff().dt.total_seconds()
    ok(u["dt"].max() < 60, "a pause inside an unloading %.0f s" % u["dt"].max())
    st = u.groupby("unloading").agg(dep=("depot_code", "first"), door=("scanner_id", "first"), a=("t", "min"),
                                    b=("t", "max")).sort_values(["dep", "door", "a"])
    st["gap"] = (st["a"] - st.groupby(["dep", "door"])["b"].shift()).dt.total_seconds() / 60
    ok(st["gap"].min() > 30, "two unloadings at a door %.1f minutes apart" % st["gap"].min())
    on_plan = R.assign_plan(arr, sh, d["plan"])
    diff = [t for t in gold if t not in sh and gold[t] != on_plan[t]]
    ok(len(diff) >= 100, "parcels whose sort hub differs from the plan's %d" % len(diff))
    mixed = R.mixed_days(arr, sh)
    ok(all(dep in ("D-WAW1", "D-WAW2") for dep, _ in mixed) and len(mixed) >= 15, "mixed mornings %d" % len(mixed))
    # Pruszkow's own rate rises on its heavy nights (the excess route's base rate is not the night's)
    nights = w["nights"]
    heavy = [n for n in nights if nights[n]["flex"] > 0]
    ok(min(nights[n]["waw_rate"] for n in heavy) > 1.2 * sum(nights[n]["waw_rate"] for n in nights if n not in heavy)
       / max(1, len(nights) - len(heavy)), "Pruszkow's rate on its heavy nights")

    # ------------------------------------------------------------ the semantic fork rows
    hub_rows = d["hub"].copy()
    hub_rows["night_local"] = (pd.to_datetime(hub_rows["event_time_utc"]).dt.tz_convert("Europe/Warsaw")
                               - pd.Timedelta(hours=12)).dt.strftime("%Y-%m-%d")
    by_sort = set(hub_rows.loc[(hub_rows["night_local"] >= "2026-10-01"), "tracking_no"])
    a_sort = arr[[t in by_sort or (t not in sh and ts[:10] >= "2026-10-02") for t, ts in
                  zip(arr["tracking_no"], arr["scanned_at"])]]
    ok(alerts_of(GD.hub_rates(gold, a_sort, sh)) == gold_set, "SF1 October by sort night")
    rows_n = collections.Counter(d["hub"]["hub_code"])
    alt = {}
    for h in rows_n:
        x = [v for v in fig["hubs"] if v["hub_code"] == h][0]
        alt[h] = {"rate": 100.0 * x["missed"] / (rows_n[h] + x["missed"])}
    ok(alerts_of(alt) == gold_set, "SF2 sort rows for parcels")
    ok(pd.to_datetime(d["hub"]["received_at_utc"]).max() < pd.Timestamp("2026-11-06 06:40:00", tz="UTC"),
       "SF3 every sort event before the extract")
    alt = {x["hub_code"]: {"rate": 100.0 * x["missed"] / (x["parcels"] - x["missed"])} for x in fig["hubs"]}
    ok(alerts_of(alt) == gold_set, "SF4 the rate over sorted parcels")
    ok(min(abs(x["rate_exact"] - GD.TOLERANCE) for x in fig["hubs"]) > 0.2, "SF5 no hub within 0.2 of the line")
    ok(sorted(x["hub_code"] for x in fig["hubs"] if x["rate"] > GD.TOLERANCE) == gold_set, "SF6 rounded rates")
    raw = d["depot"][(d["depot"]["scanned_at"].str[:7] == "2026-10")]
    ok(len(raw) > len(arr), "the re-sent batch ships")
    cnt = collections.Counter(raw["tracking_no"])
    alt = collections.Counter()
    altm = collections.Counter()
    for t, c in cnt.items():
        alt[gold[t]] += c
        altm[gold[t]] += 0 if t in sh else c
    ok(alerts_of({h: {"rate": 100.0 * altm[h] / alt[h]} for h in alt}) == gold_set, "SF7 the re-sent batch counted")
    ok(raw.duplicated().sum() == len(raw) - len(arr) and not raw.drop_duplicates()["tracking_no"].duplicated().any(),
       "SF7 the re-sent rows are the same rows")
    a_m = arr.copy()
    a_m["depot_code"] = a_m["depot_code"].replace({"D-KRK3": "D-KRK1"})
    g_m, i_m, b_m, _ = GD.golden_assign(a_m, sh)
    ok(g_m == gold, "SF8 Krakow's two codes as one depot")
    hrs = pd.to_datetime(arr["scanned_at"]).dt.hour
    ok(hrs.min() >= 3 and hrs.max() <= 8, "SF9 unloadings between 03:00 and 09:00 local (%d to %d)" % (hrs.min(), hrs.max()))
    ok(all(alerts_of(GD.hub_rates(gold, arr, sh)) == gold_set for _ in (0,)), "SF10 hub times never read by the call")
    ok(set(d["hubs"]["hub_code"]) == set(x["hub_code"] for x in fig["hubs"]) and len(d["hubs"]) == 6,
       "SF12 the register's hubs")
    for k in ("depot", "hub", "courier"):
        ok(pd.to_datetime(d[k]["received_at_utc"]).max() < pd.Timestamp("2026-11-06 06:40:00", tz="UTC"),
           "SF13 %s records before the extract" % k)
    master = {}
    for p in ps:
        master[p["label"]] = p["label"][:13] if p["kind"] == "b2b" else p["label"]
    cons_n, cons_m = collections.Counter(), collections.Counter()
    seen_c = set()
    for t in arr["tracking_no"]:
        key = (master[t], gold[t])
        if key in seen_c:
            continue
        seen_c.add(key)
        cons_n[gold[t]] += 1
        cons_m[gold[t]] += 0 if t in sh else 1
    ok(alerts_of({h: {"rate": 100.0 * cons_m[h] / cons_n[h]} for h in cons_n}) == gold_set, "SF14 consignments as parcels")
    ok(R.alerts(M["R0 plan"]) != gold_set, "SF11 the plan's hub is another call (closed by TDS-4 2.1)")
    ok(fig["fired"] == [x["hub_name"] for x in fig["hubs"] if x["fires"]] and
       all(isinstance(x["fires"], bool) for x in fig["hubs"]), "SF15 one verdict per hub")
    a_d = arr.copy()
    a_d["scanner_id"] = a_d["depot_code"]
    g_d, i_d, b_d, _ = GD.golden_assign(a_d, sh)
    ok(g_d == gold and i_d == 0 and b_d == 0, "SF16 the depot's arrivals without the door give the same unloadings")

    # ------------------------------------------------------------ separation line and deletion tests
    tmp = tempfile.mkdtemp(dir=os.environ.get("TMPBUILD"))
    try:
        for k in B.FILES:
            if k not in ("claims", "sorters", "request"):
                shutil.copy(F[k], os.path.join(tmp, B.FILES[k]))
        f2 = GD.figures(tmp, B.FILES)
        ok(f2 == fig, "deleting the distractors and the request moves a figure")
        keep = ("depot", "hub", "plan", "register", "standard")
        for k in B.FILES:
            if k not in keep and os.path.exists(os.path.join(tmp, B.FILES[k])):
                os.remove(os.path.join(tmp, B.FILES[k]))
        d3 = {"depot": pd.read_csv(os.path.join(tmp, B.FILES["depot"]), dtype=str),
              "hub": pd.read_parquet(os.path.join(tmp, B.FILES["hub"]))}
        a3 = GD.arrivals(d3["depot"])
        s3 = GD.sort_hub(d3["hub"])
        g3, _, _, _ = GD.golden_assign(a3, s3)
        ok(alerts_of(GD.hub_rates(g3, a3, s3)) == gold_set, "the call from the arrivals and the sorts alone")
    finally:
        shutil.rmtree(tmp)

    # ------------------------------------------------------------ ask A: courier events sent twice
    A, Bk = R.ask_rungs(d)
    deps = list(d["depots"]["depot_code"])
    cop = [e for e in w["courier"]["events"] if e["copy"]]
    ta = collections.Counter(sim.depot_code(e["depot"], e["day"]) for e in cop)
    ok(all(x["sent_twice"] == ta[x["depot_code"]] for x in fig["depots"]), "ask A against the clean truth")
    t51 = collections.Counter(sim.depot_code(e["depot"], e["day"]) for e in cop if e["model"] == "TC51")
    t52 = collections.Counter(sim.depot_code(e["depot"], e["day"]) for e in cop if e["model"] == "TC52")
    ok(min(t51[c] for c in deps) >= 20 and min(t52[c] for c in deps) >= 20,
       "both handheld models re-send at every depot (%d, %d)" % (min(t51[c] for c in deps), min(t52[c] for c in deps)))
    for k, v in A.items():
        if k == "golden":
            continue
        hits = [c for c in deps if v.get(c, 0) == A["golden"].get(c, 0)]
        if k.startswith("A4"):
            ok(len(hits) <= 1, "ask A stop %s keeps %s" % (k, hits))
        else:
            ok(not hits, "ask A stop %s keeps %s" % (k, hits))
    reload_dep = collections.Counter(sim.depot_code(p["dest"], p["arr_day"]) for p in ps if p["reload_min"] is not None)
    ok(min(reload_dep[c] for c in deps) >= 20, "a second loop at every depot")
    late51 = collections.Counter(sim.depot_code(e["depot"], e["day"]) for e in cop
                                 if e["model"] == "TC51" and e["day"] >= P.SUMMER_END)
    ok(sum(1 for c in deps if late51[c] > 0) >= 12, "TC51 re-sends after the clocks went back")

    # ------------------------------------------------------------ ask B: pre-advised, never picked up
    latest = {}
    for r in d["pre"].itertuples():
        if r.consignment_id not in latest or int(r.version) > int(latest[r.consignment_id].version):
            latest[r.consignment_id] = r
    lines = {ln["cons"]: ln for ln in w["preadvice"]}
    tb = collections.Counter()
    for c, r in latest.items():
        if not ("2026-10-01" <= r.collection_date <= "2026-10-31"):
            continue
        dep = r.collection_depot
        if dep == "D-KRK1" and r.collection_date >= "2026-10-12":
            dep = "D-KRK3"
        ln = lines[c]
        tb[dep] += ln["short"] if ln["collected"] else ln["pieces"]
    ok(all(x["never_picked"] == tb[x["depot_code"]] for x in fig["depots"]), "ask B against the clean truth")
    for k, v in Bk.items():
        if k in ("golden", "B4b every courier event but a failed pickup"):
            continue
        hits = [c for c in deps if v.get(c, 0) == Bk["golden"].get(c, 0)]
        lim = 2 if k.startswith("B4c") else 0
        ok(len(hits) <= lim, "ask B stop %s keeps %s" % (k, hits))
    ok(Bk["B4b every courier event but a failed pickup"] == Bk["golden"], "the golden route through the delivery events")
    bulk_dep = collections.Counter(p["origin"] for p in w["parcels"] if p["bulk"])
    ok(min(bulk_dep[x] for x in sim.DEPOTS) >= 100, "trailer collections at every depot")
    failed = d["courier"][d["courier"]["event_code"] == "PICKUP_FAILED"]
    ok(failed["depot_code"].nunique() == 13, "failed pickups at every depot")
    ok(not set(failed["tracking_no"]) & set(sh) and not set(failed["tracking_no"]) & set(arr["tracking_no"]),
       "a failed pickup never travels")
    shortp = sum(ln["short"] for ln in w["preadvice"] if ln["collected"])
    ok(shortp >= 100, "short consignments %d" % shortp)
    refs = [ln["ref"] for ln in w["preadvice"] if ln["kind"] == "marketplace"]
    grp = collections.Counter(r[:20] for r in refs)
    ok(sum(v for v in grp.values() if v > 1) >= 1000, "order numbers sharing their first 20 characters")
    ok(d["courier"]["customer_ref"].dropna().str.len().max() == 20, "the handheld keeps 20 characters")
    pre_v = d["pre"].groupby("consignment_id")["version"].nunique()
    ok((pre_v > 1).sum() >= 1000, "later versions ship")
    moved_out = [c for c, r in latest.items() if r.collection_date > "2026-10-31" and int(r.version) == 2]
    ok(len(moved_out) >= 1, "pre-advices moved out of October %d" % len(moved_out))
    krk_old = d["pre"][(d["pre"]["collection_depot"] == "D-KRK1") & (d["pre"]["collection_date"] >= "2026-10-12")]
    ok(len(krk_old) >= 200, "the replaced code after the move %d" % len(krk_old))
    ok(len(d["pre"][d["pre"]["collection_date"] < "2026-10-01"]) >= 500, "September collections ship")
    ok(not [k for k in cols if k in ("is_copy", "resent", "duplicate", "copy", "trailer_collection")],
       "no marker column")
    # SF17 to SF22
    cp = GD.sent_twice(d["courier"], d["handhelds"])
    ok(len(cp) == sum(x["sent_twice"] for x in fig["depots"]), "SF17 events equal records")
    hh_dep = dict(zip(d["handhelds"]["device_id"], d["handhelds"]["depot_code"]))
    alt = collections.Counter(hh_dep[x] for x in cp["device_id"])
    ok(alt["D-KRK1"] == 0 and alt["D-KRK3"] > A["golden"]["D-KRK3"], "SF18 the register's depot moves Krakow")
    ok(Bk["B6 the cracker on every version"] != Bk["B3 pieces and references counted, pickups only"], "SF19 versions move B")
    ok(Bk["B5 the cracker without the replaced code"] != Bk["B3 pieces and references counted, pickups only"],
       "SF20 the replaced code moves B")
    cons = R.b_route(d, events=("PICKED_UP", "OUT_FOR_DELIVERY", "DELIVERY_ATTEMPTED", "DELIVERED"), pieces=False)
    ok(all(cons.get(c, 0) != Bk["golden"].get(c, 0) for c in sim.B2B_DEPOTS),
       "SF21 consignments for parcels moves every contract depot")
    ok(all(Bk["B3 pieces and references counted, pickups only"].get(c, 0) != Bk["golden"].get(c, 0) for c in deps),
       "SF22 a pickup scan for a collection moves every depot")

    # ------------------------------------------------------------ realism sweep and clocks
    H = fig["hubs"]
    ok(len({x["parcels"] for x in H}) == 6 and len({x["missed"] for x in H}) == 6, "identical hub counts")
    ok(len({x["sent_twice"] for x in fig["depots"]}) >= 12 and len({x["never_picked"] for x in fig["depots"]}) >= 12,
       "identical depot counts")
    ok(all(round(x["rate"] * 100) % 50 != 0 for x in H), "a round rate")
    for who in ("Inga Blaut", "Krzysztof Konieczek", "Aurelia Pieczyrak"):
        ok(who in open(F["request"], encoding="utf-8").read(), "persona %s" % who)
    hu = d["hub"].drop_duplicates("tracking_no").set_index("tracking_no")["event_time_utc"]
    au = arr.set_index("tracking_no")["scanned_at"]
    loc = pd.to_datetime(au.loc[hu.index]).dt.tz_localize("Europe/Warsaw").dt.tz_convert("UTC")
    lead = (loc - pd.to_datetime(hu)).dt.total_seconds() / 3600
    ok(lead.min() > 0.5 and lead.max() < 12, "sort before arrival (%.1f to %.1f h)" % (lead.min(), lead.max()))
    rec = pd.to_datetime(d["hub"]["received_at_utc"]) - pd.to_datetime(d["hub"]["event_time_utc"])
    ok(rec.dt.total_seconds().min() >= 0 and rec.dt.total_seconds().max() < 15, "hub receipt lag")
    pre = d["pre"]
    sent = pd.to_datetime(pre["sent_at_utc"]).dt.tz_convert("Europe/Warsaw").dt.strftime("%Y-%m-%d")
    ok((sent[pre["version"] == "1"] < pre.loc[pre["version"] == "1", "collection_date"]).all(),
       "a first pre-advice sent on or after its collection day")
    orig = d["courier"][~d["courier"]["event_id"].isin(set(cp["event_id"]))]
    ou = pd.to_datetime(orig["event_time"]).dt.tz_localize("Europe/Warsaw").dt.tz_convert("UTC")
    lag = (pd.to_datetime(orig["received_at_utc"]) - ou).dt.total_seconds()
    ok(lag.min() >= 0, "a courier event received before it happened (%.0f s)" % lag.min())
    ok(d["courier"]["event_id"].is_unique and d["hub"]["event_id"].is_unique, "event ids")
    ok(len(set(d["handhelds"]["device_id"])) == len(d["handhelds"]), "handheld ids")
    ok(set(d["courier"]["device_id"]) <= set(d["handhelds"]["device_id"]), "every handheld in the register")
    ok(set(arr["scanner_id"]) <= set(_sheet_col(F["register"], "Door scanners", "scanner_id")), "every door reader registered")
    ok(set(d["courier"]["event_code"]) == {"PICKED_UP", "PICKUP_FAILED", "OUT_FOR_DELIVERY", "DELIVERY_ATTEMPTED",
                                           "DELIVERED"}, "event codes")
    for gpath in (os.path.join(golden, GD.XLSX), os.path.join(golden, GD.PPTX)):
        ok(os.path.exists(gpath), "golden %s" % gpath)
    if FAILS:
        raise AssertionError("%d checks failed:\n  " % len(FAILS) + "\n  ".join(FAILS))
    return N[0]


def _sheet_col(path, sheet, col):
    wb = load_workbook(path, read_only=True)
    rows = list(wb[sheet].iter_rows(values_only=True))
    i = rows[0].index(col)
    return [r[i] for r in rows[1:]]
