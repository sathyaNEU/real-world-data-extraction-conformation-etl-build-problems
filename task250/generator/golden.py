"""The golden figures, computed from the shipped files under target/ alone, and the two deliverables.

    fig = figures(target, FILES)      every graded figure
    write(golden_dir, fig)            tracking_scorecard_2026-10.xlsx and scan_alerts_2026-10.pptx
"""
import datetime as dt
import io
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX = "tracking_scorecard_2026-10.xlsx"
PPTX = "scan_alerts_2026-10.pptx"
MONTH = ("2026-10-01", "2026-10-31")
TOLERANCE = 1.00          # TDS-4 2.3, per cent
GAP_MIN = 10.0            # a pause at a door longer than this ends one unloading


# ------------------------------------------------------------------ reading the pack
def _sheet(path, name):
    wb = load_workbook(path, read_only=True)
    rows = list(wb[name].iter_rows(values_only=True))
    wb.close()
    return pd.DataFrame(rows[1:], columns=rows[0])


def load(target, F):
    j = lambda k: os.path.join(target, F[k])
    return {
        "depot": pd.read_csv(j("depot"), dtype=str),
        "hub": pd.read_parquet(j("hub")),
        "courier": pd.read_parquet(j("courier")),
        "pre": pd.read_csv(j("preadvice"), dtype=str, keep_default_na=False),
        "hubs": _sheet(j("register"), "Hubs"),
        "depots": _sheet(j("register"), "Depots"),
        "handhelds": _sheet(j("register"), "Handhelds"),
        "plan": _sheet(j("plan"), "Sort plan"),
    }


def arrivals(dep):
    """one arrival per parcel unloaded on an October morning (the re-sent batch's rows are the same rows again)"""
    a = dep.drop_duplicates().copy()
    a = a[(a["scanned_at"].str[:10] >= MONTH[0]) & (a["scanned_at"].str[:10] <= MONTH[1])]
    assert a["tracking_no"].is_unique, "a parcel unloaded twice"
    a["t"] = pd.to_datetime(a["scanned_at"])
    return a


def sort_hub(hub):
    n = hub.groupby("tracking_no")["hub_code"].nunique()
    assert n.max() == 1, "a parcel sorted at two hubs"
    h = hub.drop_duplicates("tracking_no")
    return dict(zip(h["tracking_no"], h["hub_code"]))


def unloadings(arr, gap=GAP_MIN):
    """each door's unloadings: arrivals at one depot door split wherever the door stood idle longer than the gap"""
    a = arr.sort_values(["depot_code", "scanner_id", "t", "tracking_no"]).copy()
    new = ((a["depot_code"] != a["depot_code"].shift()) | (a["scanner_id"] != a["scanner_id"].shift())
           | ((a["t"] - a["t"].shift()) > pd.Timedelta(minutes=gap)))
    a["unloading"] = new.cumsum()
    return a


def hub_rates(assign, arr, sh):
    """parcels handled, missed sort scans and the rate by hub, for a hub assignment of every October parcel"""
    out = {}
    for h in sorted(set(assign.values())):
        n = sum(1 for t in arr["tracking_no"] if assign[t] == h)
        m = sum(1 for t in arr["tracking_no"] if assign[t] == h and t not in sh)
        out[h] = {"parcels": n, "missed": m, "rate": 100.0 * m / n}
    return out


def golden_assign(arr, sh, gap=GAP_MIN):
    u = unloadings(arr, gap)
    assign, impure, blind = {}, 0, 0
    for _, g in u.groupby("unloading"):
        hubs = {sh[t] for t in g["tracking_no"] if t in sh}
        if len(hubs) != 1:
            impure += len(hubs) > 1
            blind += len(hubs) == 0
            continue
        h = hubs.pop()
        for t in g["tracking_no"]:
            assign[t] = h
    return assign, impure, blind, u["unloading"].nunique()


# ------------------------------------------------------------------ courier events sent twice
def to_utc_minute(local):
    t = pd.to_datetime(local).dt.tz_localize("Europe/Warsaw", ambiguous="raise", nonexistent="raise")
    return t.dt.tz_convert("UTC").dt.strftime("%Y-%m-%d %H:%M")


def sent_twice(cou, handhelds):
    """the further records of a scan the platform already held: same parcel, event, courier and handheld, the time the
    first record carries, as the handheld's re-send writes it (cut to the minute, in UTC on the TC51)"""
    model = dict(zip(handhelds["device_id"], handhelds["model"]))
    key = ["tracking_no", "event_code", "courier_id", "device_id"]
    c = cou[cou.duplicated(key, keep=False)].copy()
    c["model"] = c["device_id"].map(model)
    c["loc_min"] = c["event_time"].str[:16]
    c["utc_min"] = to_utc_minute(c["event_time"])
    c = c.sort_values(key + ["received_at_utc"])
    copies, originals = [], set()
    for _, g in c.groupby(key, sort=False):
        rows = g.to_dict("records")
        for j in range(1, len(rows)):
            b = rows[j]
            if b["event_time"][17:19] != "00":
                continue
            for a in rows[:j]:
                if a["received_at_utc"] < b["received_at_utc"] and (
                        a["loc_min"] == b["loc_min"] or (b["model"] == "TC51" and a["utc_min"] == b["loc_min"])):
                    assert a["event_id"] not in originals, "a scan sent three times"
                    originals.add(a["event_id"])
                    copies.append(b)
                    break
    df = pd.DataFrame(copies)
    return df


# ------------------------------------------------------------------ pre-advised, never picked up
def replaced_codes(depots):
    out = {}
    for r in depots.itertuples():
        if isinstance(r.replaced_by, str) and r.replaced_by:
            opened = depots.loc[depots["depot_code"] == r.replaced_by, "opened"].iloc[0]
            out[r.depot_code] = (r.replaced_by, str(opened))
    return out


def october_preadvice(pre, depots):
    p = pre.copy()
    p["v"] = p["version"].astype(int)
    p = p.sort_values(["consignment_id", "v"]).groupby("consignment_id").tail(1)
    p = p[(p["collection_date"] >= MONTH[0]) & (p["collection_date"] <= MONTH[1])].copy()
    rep = replaced_codes(depots)
    p["depot"] = [rep[c][0] if c in rep and d >= rep[c][1] else c for c, d in zip(p["collection_depot"], p["collection_date"])]
    p["pieces"] = p["pieces"].astype(int)
    return p


def never_picked(pre, depots, cou, hub, dep):
    """pre-advised October parcels Kurier Warta never collected: a parcel seen anywhere on the platform was collected;
    a consignment's pieces are counted against its piece labels, a marketplace order against the pickups carrying its
    reference (the handheld keeps the first 20 characters)"""
    p = october_preadvice(pre, depots)
    moved = cou[cou["event_code"] != "PICKUP_FAILED"]
    seen = set(moved["tracking_no"]) | set(hub["tracking_no"]) | set(dep["tracking_no"])
    pieces = {}
    for t in seen:
        if len(t) == 16 and t.isdigit():
            pieces[t[:13]] = pieces.get(t[:13], 0) + 1
    out = {}
    lab = p[p["tracking_no"] != ""]
    for r in lab.itertuples():
        got = (1 if r.tracking_no in seen else 0) + pieces.get(r.tracking_no, 0)
        assert got <= r.pieces, ("more pieces than pre-advised", r.consignment_id)
        out[r.depot] = out.get(r.depot, 0) + r.pieces - got
    mk = p[p["tracking_no"] == ""].copy()
    mk["key"] = mk["customer_ref"].str[:20]
    pu = cou[(cou["event_code"] == "PICKED_UP") & cou["customer_ref"].notna()]
    picked = pu.groupby("customer_ref")["tracking_no"].nunique().to_dict()
    for (key, depot), g in mk.groupby(["key", "depot"]):
        n = int(g["pieces"].sum())
        k = picked.get(key, 0)
        assert k <= n, ("more pickups than pre-advised orders", key)
        out[depot] = out.get(depot, 0) + n - k
    return out, p


# ------------------------------------------------------------------ everything owed
def figures(target, F):
    d = load(target, F)
    arr = arrivals(d["depot"])
    sh = sort_hub(d["hub"])
    assign, impure, blind, n_unl = golden_assign(arr, sh)
    assert impure == 0 and blind == 0 and len(assign) == len(arr), (impure, blind)
    r = hub_rates(assign, arr, sh)
    hubs = list(d["hubs"]["hub_code"])
    name = dict(zip(d["hubs"]["hub_code"], d["hubs"]["hub_name"]))
    H = []
    for h in hubs:
        x = r[h]
        H.append({"hub_code": h, "hub_name": name[h], "parcels": x["parcels"], "missed": x["missed"],
                  "rate": round(x["rate"], 2), "rate_exact": x["rate"], "fires": x["rate"] > TOLERANCE})
    cp = sent_twice(d["courier"], d["handhelds"])
    twice = cp.groupby("depot_code").size().to_dict() if len(cp) else {}
    never, _ = never_picked(d["pre"], d["depots"], d["courier"], d["hub"], d["depot"])
    dname = dict(zip(d["depots"]["depot_code"], d["depots"]["depot_name"]))
    D = []
    for c in d["depots"]["depot_code"]:
        D.append({"depot_code": c, "depot_name": dname[c], "sent_twice": int(twice.get(c, 0)),
                  "never_picked": int(never.get(c, 0))})
    fired = [x["hub_name"] for x in H if x["fires"]]
    return {"hubs": H, "depots": D, "fired": fired, "n_parcels": len(arr), "n_unloadings": n_unl,
            "tolerance": TOLERANCE, "extract": "2026-11-06 07:40"}


# ------------------------------------------------------------------ the deliverables
def _workbook(path, fig):
    wb = Workbook()
    ws = wb.active
    ws.title = "Hubs"
    bold = Font(bold=True)
    ws["A1"] = "Kurier Warta, tracking scorecard, October 2026 (TDS-4 v4.1; platform extract 6 Nov 2026 07:40)"
    ws["A1"].font = Font(bold=True, size=11)
    head = ["Hub", "Hub code", "Parcels handled", "Missed sort scans", "Missed-scan rate (%)", "Scan alert"]
    ws.append([])
    ws.append(head)
    for c in ws[3]:
        c.font = bold
        c.fill = PatternFill("solid", fgColor="DDE3EA")
    for x in fig["hubs"]:
        ws.append([x["hub_name"], x["hub_code"], x["parcels"], x["missed"], x["rate"], "Fires" if x["fires"] else "-"])
    for row in ws.iter_rows(min_row=4, max_row=3 + len(fig["hubs"])):
        row[2].number_format = "#,##0"
        row[3].number_format = "#,##0"
        row[4].number_format = "0.00"
        row[5].alignment = Alignment(horizontal="center")
    n = 4 + len(fig["hubs"])
    ws.cell(row=n + 1, column=1, value="Tolerance 1.00%; the alert fires above it (TDS-4 2.3).")
    for i, wd in enumerate([22, 10, 16, 18, 20, 11], 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = "A4"
    ws2 = wb.create_sheet("Depots")
    ws2["A1"] = "Courier events sent twice and pre-advised parcels never picked up, October 2026"
    ws2["A1"].font = Font(bold=True, size=11)
    ws2.append([])
    ws2.append(["Depot code", "Depot", "Courier events sent twice", "Pre-advised, never picked up"])
    for c in ws2[3]:
        c.font = bold
        c.fill = PatternFill("solid", fgColor="DDE3EA")
    for x in fig["depots"]:
        ws2.append([x["depot_code"], x["depot_name"], x["sent_twice"], x["never_picked"]])
    for row in ws2.iter_rows(min_row=4, max_row=3 + len(fig["depots"])):
        row[2].number_format = "#,##0"
        row[3].number_format = "#,##0"
    tot = 4 + len(fig["depots"])
    ws2.cell(row=tot, column=2, value="Network").font = bold
    ws2.cell(row=tot, column=3, value=sum(x["sent_twice"] for x in fig["depots"])).font = bold
    ws2.cell(row=tot, column=4, value=sum(x["never_picked"] for x in fig["depots"])).font = bold
    ws2.cell(row=tot + 2, column=1, value="D-KRK1 to 11 October; D-KRK3 from 12 October.")
    for i, wd in enumerate([12, 24, 24, 27], 1):
        ws2.column_dimensions[get_column_letter(i)].width = wd
    ws2.freeze_panes = "A4"
    wb.save(path)


def _chart(fig):
    H = fig["hubs"]
    names = [x["hub_name"] for x in H][::-1]
    rates = [x["rate"] for x in H][::-1]
    fires = [x["fires"] for x in H][::-1]
    f, ax = plt.subplots(figsize=(6.6, 4.1), dpi=200)
    cols = ["#B5462F" if z else "#8C96A3" for z in fires]
    ax.barh(names, rates, color=cols, height=0.62)
    ax.axvline(TOLERANCE, color="#222222", lw=1.0, ls="--")
    ax.text(TOLERANCE + 0.02, len(H) - 0.45, "tolerance 1.00%", fontsize=7.5, color="#222222", va="center")
    for i, (v, z) in enumerate(zip(rates, fires)):
        ax.text(v + 0.02, i, "%.2f%%%s" % (v, "  alert" if z else ""), va="center", fontsize=7.8,
                color="#7A2A1A" if z else "#333333", fontweight="bold" if z else "normal")
    ax.set_xlim(0, max(rates) * 1.32)
    ax.set_xlabel("Missed sort scans, % of parcels handled, October 2026", fontsize=8)
    ax.tick_params(labelsize=8)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    f.tight_layout()
    buf = io.BytesIO()
    f.savefig(buf, format="png", metadata={"Software": None})
    plt.close(f)
    buf.seek(0)
    return buf


def _slide(path, fig):
    """one board slide: the call, the chart of the six rates against the line, the hub table"""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fired = fig["fired"]
    lead = ("%s and %s" % (", ".join(fired[:-1]), fired[-1])) if len(fired) > 1 else (fired[0] if fired else "none")

    def text(x, y, w, h, runs):
        tb = s.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, (t, size, bold, col) in enumerate(runs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = "Calibri"
            r.font.color.rgb = RGBColor.from_string(col)
        return tb

    text(457200, 300000, 11277600, 900000,
         [("Scan alerts, October 2026: fire at %s" % lead, 26, True, "1F2A36"),
          ("TDS-4 v4.1, tolerance 1.00% of parcels handled. Platform extract 6 Nov 2026, 07:40.", 13, False, "4A5563")])
    s.shapes.add_picture(_chart(fig), Emu(457200), Emu(1350000), width=Emu(7000000))
    rows = [["Hub", "Parcels", "Missed", "Rate", "Alert"]]
    for x in fig["hubs"]:
        rows.append([x["hub_name"], "{:,}".format(x["parcels"]), "{:,}".format(x["missed"]), "%.2f%%" % x["rate"],
                     "fires" if x["fires"] else "-"])
    gt = s.shapes.add_table(len(rows), 5, Emu(7700000), Emu(1500000), Emu(4100000), Emu(2600000)).table
    widths = [1500000, 760000, 620000, 640000, 580000]
    for j, wd in enumerate(widths):
        gt.columns[j].width = Emu(wd)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            cell = gt.cell(i, j)
            cell.text = v
            para = cell.text_frame.paragraphs[0]
            para.runs[0].font.size = Pt(11)
            para.runs[0].font.bold = i == 0 or (j == 4 and v == "fires")
            para.runs[0].font.name = "Calibri"
    text(7700000, 4500000, 4100000, 1400000,
         [("A parcel without a sort event is charged to the hub whose trailer it came off at its delivery depot.",
           11, False, "4A5563"),
          ("Strykow sorted part of Pruszkow's flow by hand on its heavy nights and trucked it to the Warsaw depots.",
           11, False, "4A5563")])
    cp = prs.core_properties
    cp.author = "Analityka"
    cp.last_modified_by = "Analityka"
    cp.title = "Scan alerts, October 2026"
    cp.subject = "Operations board"
    cp.comments = ""
    cp.revision = 3
    prs.save(path)


def write(golden_dir, fig):
    os.makedirs(golden_dir, exist_ok=True)
    _workbook(os.path.join(golden_dir, XLSX), fig)
    _slide(os.path.join(golden_dir, PPTX), fig)
