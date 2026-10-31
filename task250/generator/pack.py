"""Cut the pack from the clean truth: the platform's three feeds (hub sort events, depot scans, courier events) as each
source writes them, the sort plan, the network register, the shippers' EDI pre-advices and accounts, and two files the
call never reads (claims, sorter availability)."""
import datetime as dt
import math
import os

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import params as P
import sim

FILES = {
    "hub": "hub_sort_events_2026-10.parquet",
    "depot": "depot_scan_events_2026-10.csv",
    "courier": "courier_events_2026-10.parquet",
    "plan": "sort_plan_2026Q4.xlsx",
    "register": "network_register_2026-10.xlsx",
    "standard": "TDS-4_tracking_data_standard.pdf",
    "preadvice": "edi_preadvice_2026-10.csv",
    "accounts": "shipper_accounts_2026-10.csv",
    "request": "october_scorecard_request.eml",
    "claims": "customer_claims_2026-10.csv",
    "sorters": "hub_sorter_availability_2026-10.csv",
}
DISTRACTORS = ["claims", "sorters"]
CONTEXT = ["request"]


# ------------------------------------------------------------------ clocks
def utc_offset_h(day, minute):
    """Europe/Warsaw: CEST (+2) until 03:00 local on 25 October 2026, CET (+1) after"""
    t = dt.datetime.combine(day, dt.time(0)) + dt.timedelta(minutes=minute)
    return 2 if t < dt.datetime(2026, 10, 25, 3, 0) else 1


def local_dt(day, minute):
    return dt.datetime.combine(day, dt.time(0)) + dt.timedelta(minutes=minute)


def to_utc(day, minute):
    return local_dt(day, minute) - dt.timedelta(hours=utc_offset_h(day, minute))


def fmt_local(t):
    return t.strftime("%Y-%m-%d %H:%M:%S")


def fmt_utc(t):
    return t.strftime("%Y-%m-%dT%H:%M:%SZ")


def uid(g):
    b = g.integers(0, 256, size=16, dtype=np.uint8)
    b[6] = (b[6] & 0x0F) | 0x40
    b[8] = (b[8] & 0x3F) | 0x80
    h = "".join("%02x" % int(x) for x in b)
    return "%s-%s-%s-%s-%s" % (h[:8], h[8:12], h[12:16], h[16:20], h[20:])


# ------------------------------------------------------------------ the feeds
def hub_feed(w, seed=P.SEED):
    """one row per read the hub's interface sent: the sort, and a second read where the parcel went round the loop"""
    g = np.random.default_rng([seed, 1, 11])
    scanners = {"HPOZ": ["POZ-T1", "POZ-T2", "POZ-H07"], "HWRO": ["WRO-T1", "WRO-T2", "WRO-H03"],
                "HKTW": ["KTW-T1", "KTW-H05"], "HKRK": ["KRK-T1", "KRK-T2", "KRK-H02"],
                "HLDZ": ["LDZ-T1", "LDZ-T2", "LDZ-H11", "LDZ-H12", "LDZ-H14"],
                "HWAW": ["WAW-T1", "WAW-T2", "WAW-T3", "WAW-H04"]}
    rows = []
    for p in sorted(w["parcels"], key=lambda p: p["pid"]):
        if not p["scanned"]:
            continue
        hub = p["hub"]
        if p["flex"]:
            sc = scanners["HLDZ"][2 + p["pid"] % 3]          # the manual line's handhelds
        else:
            sc = scanners[hub][int(g.integers(0, len(scanners[hub]) - 1))] if g.random() < 0.97 else scanners[hub][-1]
        t = to_utc(p["night"], p["sort_min"])
        rows.append((t, p["label"], hub, sc))
        if p["recirc"]:
            rows.append((t + dt.timedelta(seconds=float(g.uniform(160, 540))), p["label"], hub, sc))
    rows.sort(key=lambda r: (r[0], r[1]))
    out = []
    for t, lab, hub, sc in rows:
        rec = t + dt.timedelta(seconds=float(g.uniform(0.4, 9.0)))
        out.append({"event_id": uid(g), "tracking_no": lab, "hub_code": hub, "scanner_id": sc,
                    "event_time_utc": fmt_utc(t), "received_at_utc": fmt_utc(rec)})
    return pd.DataFrame(out)


def depot_feed(w, seed=P.SEED):
    """one row per arrival scan at a depot door, as the depot server sends it; the server re-sent one morning's
    batch at two depots on 21 October with the rows as they were"""
    g = np.random.default_rng([seed, 1, 12])
    rows = []
    for p in w["parcels"]:
        t = local_dt(p["arr_day"], p["arr_min"])
        rows.append((t, p["label"], p["depot_code"], p["door"], p["arr_day"], p["arr_min"]))
    rows.sort(key=lambda r: (r[0], r[2], r[1]))
    out = []
    for t, lab, code, door, day, minute in rows:
        rec = to_utc(day, minute) + dt.timedelta(seconds=float(g.uniform(1, 25)))
        out.append({"event_id": uid(g), "tracking_no": lab, "depot_code": code, "scanner_id": door,
                    "event_type": "ARRIVED", "scanned_at": fmt_local(t), "received_at_utc": fmt_utc(rec)})
    df = pd.DataFrame(out)
    resent = df[(df["scanned_at"].str.startswith("2026-10-21")) & (df["depot_code"].isin(["D-WRO1", "D-KTW1"]))]
    resent = resent[resent["scanned_at"].str[11:13].isin(["04", "05"])]
    df = pd.concat([df, resent], ignore_index=True)
    return df


def courier_feed(w, seed=P.SEED):
    """the couriers' handheld events; a re-send carries the event as the app queued it (TC52: local time, TC51: UTC,
    both cut to the minute) and goes out on a later sync"""
    g = np.random.default_rng([seed, 1, 13])
    names = {"PICKUP": "PICKED_UP", "PICKUP_FAILED": "PICKUP_FAILED", "OUT_FOR_DELIVERY": "OUT_FOR_DELIVERY",
             "ATTEMPT": "DELIVERY_ATTEMPTED", "DELIVERED": "DELIVERED"}
    rows = []
    for e in w["courier"]["events"]:
        t = local_dt(e["day"], e["min"])
        tu = to_utc(e["day"], e["min"])
        rec = tu + dt.timedelta(minutes=e["sync_lag"]) + dt.timedelta(seconds=float(g.uniform(1, 40)))
        if e["copy"]:
            if e["model"] == "TC51":
                shown = tu.replace(second=0)
            else:
                shown = t.replace(second=0)
        else:
            shown = t
        code = sim.depot_code(e["depot"], e["day"])
        rows.append((rec, {"event_id": None, "tracking_no": e["label"], "depot_code": code, "courier_id": e["courier"],
                           "device_id": e["device"], "event_code": names[e["type"]], "event_time": fmt_local(shown),
                           "received_at_utc": fmt_utc(rec), "customer_ref": e["ref"][:20] if e["ref"] else None}))
    rows.sort(key=lambda r: (r[0], r[1]["tracking_no"], r[1]["event_code"]))
    out = []
    for rec, r in rows:
        r["event_id"] = uid(g)
        out.append(r)
    return pd.DataFrame(out)


def preadvice_file(w, seed=P.SEED):
    """the EDI pre-advice lines as received: one line per consignment, a later version where a shipper corrected one
    (a weight, or the collection day moved)"""
    g = np.random.default_rng([seed, 1, 14])
    out = []
    for ln in w["preadvice"]:
        day = ln["date"]
        sent = to_utc(day - dt.timedelta(days=1), float(g.uniform(15 * 60, 23.5 * 60)))
        depot = ln["depot"]
        if depot == "D-KRK1" and day >= P.KRK_MOVE and g.random() < 0.55:
            dep_code = "D-KRK1"                        # shippers' systems still carrying the old code
        else:
            dep_code = sim.depot_code(depot, day)
        base = {"consignment_id": ln["cons"], "version": 1, "sent_at_utc": fmt_utc(sent),
                "shipper_account": ln["shipper"], "collection_depot": dep_code,
                "collection_date": day.isoformat(), "tracking_no": ln["master"] or None,
                "pieces": ln["pieces"], "customer_ref": ln["ref"] or None,
                "weight_kg": round(float(g.lognormal(1.0, 0.7)) * ln["pieces"], 1)}
        out.append(base)
        if g.random() < 0.06:
            v2 = dict(base, version=2, sent_at_utc=fmt_utc(sent + dt.timedelta(hours=float(g.uniform(1, 14)))))
            if not ln["collected"] and day.weekday() == 4 and day >= dt.date(2026, 10, 23) and g.random() < 0.7:
                v2["collection_date"] = (day + dt.timedelta(days=3)).isoformat()   # moved to the next Monday
            else:
                v2["weight_kg"] = round(base["weight_kg"] * float(g.uniform(0.7, 1.3)), 1)
            out.append(v2)
    out.sort(key=lambda r: (r["sent_at_utc"], r["consignment_id"], r["version"]))
    return pd.DataFrame(out)


def accounts_file(w, seed=P.SEED):
    g = np.random.default_rng([seed, 1, 15])
    stems = ["Bialski", "Gorecki", "Kaminski", "Lesniak", "Malinowski", "Nowicki", "Ostrowski", "Pawlak", "Rutkowski",
             "Sadowski", "Tomczak", "Wawrzyniak", "Zielinski", "Adamczyk", "Borowski", "Czarnecki", "Dudek", "Jasinski",
             "Kubiak", "Laskowski", "Mazur", "Pietrzak", "Sikora", "Urban", "Wieczorek", "Wrobel", "Zawadzki"]
    trades = ["Hurtownia", "Drukarnia", "Sklep", "Zaklad", "Pracownia", "Studio", "Apteka", "Salon", "Magazyn"]
    goods = ["Elektro", "Agro", "Medi", "Tex", "Moto", "Pol", "Bud", "Eko", "Opti", "Libra", "Nova", "Vita"]
    integ = {"api": "API", "marketplace": "MARKETPLACE", "b2b": "EDI-B2B", "legacy": "EDI-LEGACY"}
    out = []
    for s in w["shippers"]:
        r = g.random()
        if r < 0.34:
            name = "%s %s sp. z o.o." % (goods[int(g.integers(0, len(goods)))] + ["tron", "mex", "plast", "trade", "med"][int(g.integers(0, 5))], ["", "Polska"][int(g.integers(0, 2))])
            name = " ".join(name.split())
        elif r < 0.6:
            name = "%s %s" % (trades[int(g.integers(0, len(trades)))], stems[int(g.integers(0, len(stems)))])
        elif r < 0.8:
            name = "PPHU %s %s" % (stems[int(g.integers(0, len(stems)))], "i Syn" if g.random() < 0.3 else "")
            name = name.strip()
        else:
            name = "%s S.A." % (goods[int(g.integers(0, len(goods)))] + ["pak", "dom", "lux", "flex"][int(g.integers(0, 4))])
        start = dt.date(2017, 1, 1) + dt.timedelta(days=int(g.integers(0, 3300)))
        out.append({"shipper_account": s["account"], "shipper_name": name, "collection_depot": sim.depot_code(s["depot"], dt.date(2026, 10, 31)),
                    "integration": integ[s["kind"]], "contract_start": start.isoformat()})
    return pd.DataFrame(out)


def claims_file(w, seed=P.SEED):
    g = np.random.default_rng([seed, 1, 16])
    ps = w["parcels"]
    out = []
    cid = 2600000 + int(g.integers(0, 50000))
    for p in ps:
        if g.random() < 0.0095:
            cid += int(g.integers(1, 12))
            kind = ["DAMAGED", "LATE", "CONTENTS_MISSING", "LOST"][int(g.choice(4, p=[0.46, 0.31, 0.15, 0.08]))]
            opened = p["deliv_day"] + dt.timedelta(days=int(g.integers(0, 9)))
            amount = round(float(g.lognormal(4.6, 0.9)), 2)
            status = ["OPEN", "ACCEPTED", "REJECTED", "PAID"][int(g.integers(0, 4))]
            if opened > P.EXTRACT.date():
                continue                               # not yet opened when the extract was taken
            if (P.EXTRACT.date() - opened).days < 5:
                status = "OPEN"
            out.append({"claim_id": "RK%07d" % cid, "tracking_no": p["label"], "claim_type": kind,
                        "opened_on": opened.isoformat(), "claimed_pln": amount, "status": status})
    return pd.DataFrame(out)


def sorters_file(w, seed=P.SEED):
    g = np.random.default_rng([seed, 1, 17])
    out = []
    for night in sorted(w["nights"]):
        for h in [x[0] for x in P.HUBS]:
            sched = int(g.integers(255, 300))
            down = int(max(0, g.normal(9, 8))) if g.random() < 0.55 else 0
            reason = "" if down == 0 else ["jam", "induction fault", "PLC restart", "scale calibration", "chute full"][int(g.integers(0, 5))]
            out.append({"hub_code": h, "sort_night": night.isoformat(), "scheduled_minutes": sched,
                        "downtime_minutes": down, "main_reason": reason or None})
    return pd.DataFrame(out)


# ------------------------------------------------------------------ the workbooks
def _sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(list(r))
    for i, wd in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = "A2"


def plan_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sort plan"
    rows = []
    for d in P.DEPOTS:
        code, name, hub = d[0], d[1], d[2]
        if code == "D-KRK1":
            rows.append((code, name, hub, "2026-09-28", "2026-10-11", "Q4-1"))
            rows.append((P.KRK_NEW[0], P.KRK_NEW[1], hub, "2026-10-12", "2026-12-31", "Q4-1"))
        elif code == "D-WRO2":
            rows.append((code, name, hub, "2026-09-28", "2026-10-18", "Q4-1"))
            rows.append((code, name, "HKTW", "2026-10-19", "2026-12-31", "Q4-2"))
        else:
            rows.append((code, name, hub, "2026-09-28", "2026-12-31", "Q4-1"))
    _sheet(ws, ["delivery_depot", "depot_name", "sort_hub", "first_sort_night", "last_sort_night", "plan_issue"],
           rows, [15, 24, 10, 16, 16, 11])
    ws2 = wb.create_sheet("Issues")
    _sheet(ws2, ["plan_issue", "issued", "change"],
           [("Q4-1", "2026-09-22", "Q4 plan from the sort night of 28 September; Krakow-Zablocie (D-KRK1) moves to "
                                   "Krakow-Rybitwy (D-KRK3) from 12 October"),
            ("Q4-2", "2026-10-14", "Opole (D-WRO2) sorted at Sosnowiec from the night of 19 October")],
           [11, 12, 86])
    wb.save(path)


def register_xlsx(path, w, seed=P.SEED):
    g = np.random.default_rng([seed, 1, 18])
    wb = Workbook()
    ws = wb.active
    ws.title = "Hubs"
    addr = {"HPOZ": "ul. Poznanska 160, 62-052 Komorniki", "HWRO": "ul. Logistyczna 6, 55-040 Bielany Wroclawskie",
            "HKTW": "ul. Baczynskiego 14, 41-203 Sosnowiec", "HKRK": "ul. Krakowska 29, 32-085 Modlniczka",
            "HLDZ": "ul. Gen. Sikorskiego 7, 95-010 Strykow", "HWAW": "ul. Przejazdowa 21, 05-800 Pruszkow"}
    sorter = {"HPOZ": "cross-belt, 2 induction lines", "HWRO": "cross-belt, 2 induction lines",
              "HKTW": "tilt-tray, 1 induction line", "HKRK": "cross-belt, 2 induction lines",
              "HLDZ": "cross-belt, 2 induction lines; manual sort hall", "HWAW": "cross-belt, 3 induction lines"}
    _sheet(ws, ["hub_code", "hub_name", "address", "sorter"],
           [(h[0], h[1], addr[h[0]], sorter[h[0]]) for h in P.HUBS], [10, 22, 46, 46])
    ws = wb.create_sheet("Depots")
    daddr = {"D-POZ1": "ul. Gluszyna 120, 61-329 Poznan", "D-POZ2": "ul. Okrezna 33, 64-100 Leszno",
             "D-WRO1": "ul. Stabłowicka 147, 54-066 Wroclaw".replace("ł", "l"), "D-WRO2": "ul. Wspolna 22, 45-837 Opole",
             "D-KTW1": "ul. Pukowca 5, 40-847 Katowice", "D-KTW2": "ul. Lesna 11, 43-300 Bielsko-Biala",
             "D-KRK1": "ul. Przemyslowa 3, 30-701 Krakow", "D-KRK2": "ul. Komunalna 2, 33-100 Tarnow",
             "D-LDZ1": "ul. Malownicza 4, 92-702 Lodz", "D-LDZ2": "ul. Wronia 65, 97-300 Piotrkow Trybunalski",
             "D-WAW1": "ul. Komitetu Obrony Robotnikow 45, 02-146 Warszawa",
             "D-WAW2": "ul. Handlowa 5, 03-216 Warszawa"}
    rows = []
    for d in P.DEPOTS:
        if d[0] == "D-KRK1":
            rows.append((d[0], d[1], daddr[d[0]], "2014-03-03", "2026-10-11", P.KRK_NEW[0]))
            rows.append((P.KRK_NEW[0], P.KRK_NEW[1], "ul. Rybitwy 104, 30-722 Krakow", "2026-10-12", None, None))
        else:
            rows.append((d[0], d[1], daddr[d[0]], ["2012-05-14", "2015-09-01", "2018-02-19", "2019-06-03"][int(g.integers(0, 4))], None, None))
    _sheet(ws, ["depot_code", "depot_name", "address", "opened", "closed", "replaced_by"], rows, [12, 24, 50, 12, 12, 12])
    ws = wb.create_sheet("Door scanners")
    rows = []
    for d in P.DEPOTS:
        codes = [d[0]] if d[0] != "D-KRK1" else ["D-KRK1", P.KRK_NEW[0]]
        for c in codes:
            for door in ("B1", "B2"):
                inst = "2026-10-12" if c == P.KRK_NEW[0] else ["2021-04-12", "2022-08-29", "2023-11-06"][int(g.integers(0, 3))]
                rows.append((c, door, "%s-%s" % (c, door), "fixed tunnel reader", inst))
    _sheet(ws, ["depot_code", "door", "scanner_id", "scanner_type", "installed"], rows, [12, 7, 14, 20, 12])
    ws = wb.create_sheet("Handhelds")
    rows = []
    for d, cs in w["courier"]["couriers"].items():
        for c in cs:
            issued = dt.date(2021, 1, 4) + dt.timedelta(days=int(g.integers(0, 1700)))
            if c["model"] == "TC51":
                issued = dt.date(2019, 3, 4) + dt.timedelta(days=int(g.integers(0, 700)))
            rows.append((c["device"], sim.depot_code(d, dt.date(2026, 10, 31)), c["courier"], c["model"], issued.isoformat()))
    _sheet(ws, ["device_id", "depot_code", "courier_id", "model", "issued"], rows, [12, 12, 12, 8, 12])
    wb.save(path)


def write_tables(w, out):
    f = lambda k: os.path.join(out, FILES[k])
    hub = hub_feed(w)
    hub.to_parquet(f("hub"), index=False, compression="zstd")
    dep = depot_feed(w)
    dep.to_csv(f("depot"), index=False)
    cou = courier_feed(w)
    cou.to_parquet(f("courier"), index=False, compression="zstd")
    pre = preadvice_file(w)
    pre.to_csv(f("preadvice"), index=False)
    accounts_file(w).to_csv(f("accounts"), index=False)
    claims_file(w).to_csv(f("claims"), index=False)
    sorters_file(w).to_csv(f("sorters"), index=False)
    plan_xlsx(f("plan"))
    register_xlsx(f("register"), w)
    return {"hub": len(hub), "depot": len(dep), "courier": len(cou), "preadvice": len(pre)}
