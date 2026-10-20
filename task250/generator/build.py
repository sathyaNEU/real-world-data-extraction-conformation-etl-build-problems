"""Write the pack: every file under target/, from the seeded world.   B.run(target) -> the world and the record"""
import datetime as dt
import os

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import docs
import params as P
import postprocess
import records as RC
import sim

SEED = 1
FILES = dict(
    daily="estaciones_diario_2025-2027.csv",
    qc="control_calidad_2025-2027.parquet",
    register="registro_estaciones.xlsx",
    procedure="QC-02_control_calidad_temperatura_v4.pdf",
    orders="ordenes_trabajo_2025-2027.xlsx",
    checks="comprobaciones_patron_2025-2027.csv",
    request="configuracion_alertas_2028.eml",
    plan="plan_campo_2028.pdf",
    summary="resumen_alertas_2027.pdf",
    dictionary="data_dictionary.md",
    delivery="entregas_plataforma_2027.csv",
    et0="boletin_et0_2027.csv",
)
PACK = list(FILES)
DISTRACTORS = ["delivery", "et0"]
CONTEXT = ["summary"]
STAMP_PDF = {"procedure": "2024-11-18T10:05:00", "plan": "2027-12-20T12:31:00", "summary": "2028-01-02T07:15:00"}


def world():
    W = sim.world(SEED)
    rec = RC.recorded(W, SEED)
    reg = RC.register(W, SEED)
    return W, rec, reg


def _orders_xlsx(path, old, new):
    wb = Workbook()
    ws = wb.active
    ws.title = P.OLD_CONTRACTOR.split()[2]
    ws.append(list(old.columns))
    for r in old.itertuples(index=False):
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in (row[3], row[4]):
            if c.value is not None:
                c.number_format = "dd/mm/yyyy"
    for j, w in enumerate([13, 9, 15, 11, 11, 9, 30], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws2 = wb.create_sheet(P.NEW_CONTRACTOR.split()[2])
    ws2.append(list(new.columns))
    for r in new.itertuples(index=False):
        ws2.append(list(r))
    for row in ws2.iter_rows(min_row=2):
        for c in (row[3], row[4]):
            if c.value is not None:
                c.number_format = "yyyy-mm-dd"
    for j, w in enumerate([12, 16, 12, 11, 11, 10, 18], 1):
        ws2.column_dimensions[get_column_letter(j)].width = w
    for s in (ws, ws2):
        for c in s[1]:
            c.font = Font(bold=True)
        s.freeze_panes = "A2"
    wb.save(path)


def _delivery(W, reg, rng):
    rows = []
    for i, r in reg.iterrows():
        for k in range((dt.date(2027, 1, 1) - P.START).days, P.DAYS):
            if not W["avail"][k, i]:
                continue
            d = RC.day(k)
            hh = rng.choice([0, 0, 0, 0, 1, 1, 2, 5, 9]) if rng.random() < 0.2 else 0
            t = dt.datetime.combine(d + dt.timedelta(days=1), dt.time(hh, int(rng.integers(0, 60)), int(rng.integers(0, 60))))
            rows.append((r["code"], d.strftime("%d/%m/%Y"), t.strftime("%d/%m/%Y %H:%M:%S")))
    return pd.DataFrame(rows, columns=["estacion", "dia", "recibido"])


def _et0(W, rec, reg):
    rows = []
    names = dict(P.DISTRICTS)
    T = rec["T"]
    for c, nm in P.DISTRICTS:
        idx = np.where(reg["district"] == c)[0]
        for w in range(52):
            k0 = (dt.date(2027, 1, 4) - P.START).days + 7 * w
            tm = float(np.mean(T[k0:k0 + 7, idx]))
            et0 = max(0.4, 0.0023 * (tm + 17.8) * 11 * 3.3 * (1.2 + 0.8 * np.sin(2 * np.pi * (w - 12) / 52)))
            rows.append((nm, (RC.day(k0)).strftime("%d/%m/%Y"), round(tm, 1), round(et0 * 7, 1)))
    return pd.DataFrame(rows, columns=["comarca", "semana_desde", "tmed_semana", "et0_mm_semana"])


def run(target):
    os.makedirs(target, exist_ok=True)
    F = {k: os.path.join(target, v) for k, v in FILES.items()}
    W, rec, reg = world()
    daily = RC.daily_export(W, rec, reg, SEED)
    qcx = RC.qc_export(W, rec, reg)
    q2 = qcx.copy()
    q2["fecha"] = pd.to_datetime(q2["fecha"])
    q2.to_parquet(F["qc"], index=False)
    visits, checks, truth = RC.field_records(W, rec, reg, SEED, daily)
    daily = RC.complete_daily(daily, checks, reg, rec, SEED)
    d2 = daily.copy()
    d2["fecha"] = [d.strftime("%d/%m/%Y") for d in d2["fecha"]]
    d2.to_csv(F["daily"], sep=";", decimal=",", index=False, encoding="utf-8")
    old, new, ck, _ = RC.orders_and_checks(W, rec, reg, visits, checks, truth, SEED)
    _orders_xlsx(F["orders"], old, new)
    ck.to_csv(F["checks"], index=False, encoding="utf-8")
    hist = RC.adjustments(W, rec, reg, truth, SEED)
    docs.register_xlsx(F["register"], reg, hist)
    docs.procedure(F["procedure"], FILES)
    closing = set(W["closing"])
    n_close = len(closing)
    n_taken = P.N_PROGRAMME - n_close
    bd = [t for t in truth if t["kind"] == "breakdown" and t["status"] == "visit" and t["v"] is not None
          and RC.day(t["v"]).year in (2025, 2026) and t["i"] not in closing]
    docs.field_plan(F["plan"], FILES, len(reg) - n_close, n_close, n_taken, int(round(len(bd) / 2.0)))
    per_year = int(5 * round(len(rec["alerts"]) / 3.0 / 5))
    docs.request(F["request"], FILES, n_close, n_taken, per_year)
    docs.alert_summary(F["summary"], FILES, qcx, reg)
    docs.dictionary(F["dictionary"], FILES)
    rng = np.random.default_rng(SEED + 60)
    _delivery(W, reg, rng).to_csv(F["delivery"], index=False, encoding="utf-8")
    _et0(W, rec, reg).to_csv(F["et0"], index=False, encoding="utf-8", sep=";", decimal=",")
    for k in ("procedure", "plan", "summary"):
        postprocess.scrub(F[k], "Microsoft Word 2019", STAMP_PDF[k])
    for k in ("register", "orders"):
        postprocess.scrub(F[k], P.ORG, "2028-01-04")
        postprocess.finish_office(F[k], "2028-01-04T08:10:00Z", "2028-01-04T08:12:00Z", P.ORG)
    return {"W": W, "rec": rec, "reg": reg, "truth": truth, "daily": daily, "qcx": qcx, "hist": hist,
            "miskeyed": [t for t in truth if t.get("miskeyed")], "checks": checks}
