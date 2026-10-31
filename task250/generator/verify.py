"""The independent verifier: every graded figure recomputed from a target folder by a second code path that shares
nothing with golden.py (each unscanned parcel's hub from its nearest sorted neighbours at the same depot door; the
re-sends with a hand-written Warsaw clock; the pickups through merges), compared with figures.json.

    python3 verify.py [target]      -> exit 1 on any mismatch
"""
import collections
import datetime as dt
import json
import os
import sys

import pandas as pd
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
NAMES = {"depot": "depot_scan_events_2026-10.csv", "hub": "hub_sort_events_2026-10.parquet",
         "courier": "courier_events_2026-10.parquet", "pre": "edi_preadvice_2026-10.csv",
         "register": "network_register_2026-10.xlsx"}
WINDOW = dt.timedelta(minutes=15)
SUMMER_END = dt.datetime(2026, 10, 25, 3, 0)


def sheet(path, name):
    rows = list(load_workbook(path, read_only=True)[name].iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:]]


def hubs_by_neighbour(dep, hub):
    """an unscanned parcel takes the hub of the sorted parcels scanned off at the same door within 15 minutes"""
    seen = set()
    arr = []
    for r in dep.itertuples(index=False):
        k = (r.event_id, r.tracking_no)
        if k in seen:
            continue
        seen.add(k)
        if r.scanned_at.startswith("2026-10"):
            arr.append((r.depot_code, r.scanner_id, dt.datetime.strptime(r.scanned_at, "%Y-%m-%d %H:%M:%S"),
                        r.tracking_no))
    hub_of = {}
    for t, h in zip(hub["tracking_no"], hub["hub_code"]):
        assert hub_of.setdefault(t, h) == h
    lanes = collections.defaultdict(list)
    for d, door, t, tr in arr:
        lanes[(d, door)].append((t, tr))
    out = {}
    for lane, v in lanes.items():
        v.sort()
        for i, (t, tr) in enumerate(v):
            if tr in hub_of:
                out[tr] = hub_of[tr]
                continue
            near = set()
            for j in range(i - 1, -1, -1):
                if t - v[j][0] > WINDOW:
                    break
                if v[j][1] in hub_of:
                    near.add(hub_of[v[j][1]])
            for j in range(i + 1, len(v)):
                if v[j][0] - t > WINDOW:
                    break
                if v[j][1] in hub_of:
                    near.add(hub_of[v[j][1]])
            assert len(near) == 1, (lane, t, near)
            out[tr] = near.pop()
    return out, hub_of


def warsaw_to_utc(s):
    t = dt.datetime.strptime(s, "%Y-%m-%d %H:%M:%S")
    return t - dt.timedelta(hours=2 if t < SUMMER_END else 1)


def resends(cou, models):
    groups = collections.defaultdict(list)
    for r in cou.itertuples(index=False):
        groups[(r.tracking_no, r.event_code, r.courier_id, r.device_id)].append(r)
    out = collections.Counter()
    for k, v in groups.items():
        if len(v) < 2:
            continue
        v.sort(key=lambda r: r.received_at_utc)
        for j, b in enumerate(v):
            if not b.event_time.endswith(":00"):
                continue
            for a in v[:j]:
                same_local = a.event_time[:16] == b.event_time[:16]
                same_utc = models[b.device_id] == "TC51" and warsaw_to_utc(a.event_time).strftime("%Y-%m-%d %H:%M") == b.event_time[:16]
                if same_local or same_utc:
                    out[b.depot_code] += 1
                    break
    return out


def never_collected(pre, cou, hub, dep, depots):
    pre = pre.copy()
    pre["version"] = pre["version"].astype(int)
    last = pre.loc[pre.groupby("consignment_id")["version"].idxmax()]
    last = last[last["collection_date"].between("2026-10-01", "2026-10-31")].copy()
    moved = {r["depot_code"]: (r["replaced_by"], None) for r in depots if r["replaced_by"]}
    opened = {r["depot_code"]: str(r["opened"])[:10] for r in depots}
    last["at"] = [moved[c][0] if c in moved and day >= opened[moved[c][0]] else c
                  for c, day in zip(last["collection_depot"], last["collection_date"])]
    last["pieces"] = last["pieces"].astype(int)
    travelled = pd.concat([cou.loc[cou["event_code"].isin(["PICKED_UP", "OUT_FOR_DELIVERY", "DELIVERY_ATTEMPTED",
                                                            "DELIVERED"]), "tracking_no"],
                           hub["tracking_no"], dep["tracking_no"]]).drop_duplicates()
    lab = pd.DataFrame({"t": travelled})
    lab["master"] = lab["t"].where(~lab["t"].str.fullmatch(r"\d{16}"), lab["t"].str[:13])
    got = lab.groupby("master").size()
    num = last[last["tracking_no"] != ""].copy()
    num["got"] = num["tracking_no"].map(got).fillna(0).astype(int)
    assert (num["got"] <= num["pieces"]).all()
    num["left"] = num["pieces"] - num["got"]
    out = num.groupby("at")["left"].sum().to_dict()
    mk = last[last["tracking_no"] == ""].copy()
    mk["k"] = mk["customer_ref"].str.slice(0, 20)
    pk = cou[(cou["event_code"] == "PICKED_UP") & cou["customer_ref"].notna()].drop_duplicates(["tracking_no"])
    picked = pk.groupby("customer_ref").size()
    g = mk.groupby(["k", "at"])["pieces"].sum().reset_index()
    g["p"] = g["k"].map(picked).fillna(0).astype(int)
    assert (g["p"] <= g["pieces"]).all()
    g["left"] = g["pieces"] - g["p"]
    for at, v in g.groupby("at")["left"].sum().items():
        out[at] = out.get(at, 0) + int(v)
    return out


def main(target=None):
    target = target or os.path.join(TASK, "target")
    p = lambda k: os.path.join(target, NAMES[k])
    dep = pd.read_csv(p("depot"), dtype=str)
    hub = pd.read_parquet(p("hub"))
    cou = pd.read_parquet(p("courier"))
    pre = pd.read_csv(p("pre"), dtype=str, keep_default_na=False)
    reg_h = sheet(p("register"), "Hubs")
    reg_d = sheet(p("register"), "Depots")
    models = {r["device_id"]: r["model"] for r in sheet(p("register"), "Handhelds")}
    fig = json.load(open(os.path.join(HERE, "figures.json")))
    fails = []
    hub_of_parcel, _ = hubs_by_neighbour(dep, hub)
    n, m = collections.Counter(), collections.Counter()
    sorted_set = set(hub["tracking_no"])
    for t, h in hub_of_parcel.items():
        n[h] += 1
        m[h] += t not in sorted_set
    for x in fig["hubs"]:
        h = x["hub_code"]
        rate = round(100.0 * m[h] / n[h], 2)
        for k, v in (("parcels", n[h]), ("missed", m[h]), ("rate", rate), ("fires", 100.0 * m[h] / n[h] > 1.0)):
            if x[k] != v:
                fails.append("%s %s: %s against %s" % (h, k, x[k], v))
    rs = resends(cou, models)
    nc = never_collected(pre, cou, hub, dep, reg_d)
    for x in fig["depots"]:
        c = x["depot_code"]
        if x["sent_twice"] != rs.get(c, 0):
            fails.append("%s sent twice: %s against %s" % (c, x["sent_twice"], rs.get(c, 0)))
        if x["never_picked"] != nc.get(c, 0):
            fails.append("%s never picked up: %s against %s" % (c, x["never_picked"], nc.get(c, 0)))
    fired = sorted(x["hub_name"] for x in fig["hubs"] if x["fires"])
    if sorted(fig["call"]) != fired:
        fails.append("call %s against %s" % (fig["call"], fired))
    if len(reg_h) != len(fig["hubs"]):
        fails.append("hubs %d against %d" % (len(reg_h), len(fig["hubs"])))
    return fails


if __name__ == "__main__":
    f = main(sys.argv[1] if len(sys.argv) > 1 else None)
    print("verify: %s" % ("PASS" if not f else "FAIL\n  " + "\n  ".join(f)))
    sys.exit(1 if f else 0)
