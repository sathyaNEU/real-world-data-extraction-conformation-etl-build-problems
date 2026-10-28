"""An independent re-read of the shipped pack and the goldens: the import file recounted, the workbook and the memo
matched to figures.json."""
from __future__ import annotations

import json
import os
import sys

import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
T = os.path.join(TASK, "target")
GD = os.path.join(TASK, "golden")
sys.path.insert(0, HERE)
import params as P  # noqa: E402

fig = json.load(open(os.path.join(HERE, "figures.json")))
res = []


def check(label, ok, detail=""):
    res.append((label, ok, detail))


imp = pd.read_csv(os.path.join(GD, "bank_import_2027-28.csv"), dtype={"enemy_set": str}).fillna("")
bank = pd.read_csv(os.path.join(T, "tessaline_item_bank_export_20270331.csv"))
for g in P.GRADES:
    sub = imp[imp.grade == g]
    low = int((sub.difficulty < P.CUT_QSS[g]).sum())
    check("V1 grade %d count below the cut from the import file" % g, low == fig["low_%d" % g], "%d" % low)
    check("V2 grade %d commission" % g, max(0, P.FLOOR - low) == fig["commission_%d" % g], "")
    sets = sub.enemy_set[sub.enemy_set != ""]
    check("V3 grade %d enemy sets" % g, (sets.nunique(), len(sets)) == (fig["enemy_sets_%d" % g], fig["enemy_items_%d" % g]), "")
    for c in sub.reporting_domain.unique():
        check("V4 grade %d domain %s" % (g, c), int((sub.reporting_domain == c).sum()) == fig["dom_%d_%s" % (g, c)], "")
launch = bank[bank.difficulty_basis == "launch calibration 2021"].set_index("item_key").difficulty_qss
both = imp.set_index("item_key").difficulty.reindex(launch.index).dropna()
check("V5 launch items carry their bank values", (abs(both - launch.reindex(both.index)) < 1e-9).all(), "")
wk = bank[bank.difficulty_basis == "writing panel rating"].set_index("item_key").difficulty_qss
moved = imp.set_index("item_key").difficulty.reindex(wk.index)
check("V6 every workshop item carries a calibrated value, not its rating", moved.notna().all() and (moved != wk).mean() > 0.95, "")
check("V7 the import has one row per item", imp.item_key.is_unique and len(imp) == fig["import_rows"], "%d" % len(imp))
import openpyxl  # noqa: E402
wb = openpyxl.load_workbook(os.path.join(GD, "pool_summary_2027-28.xlsx"), data_only=True)
ws = wb["Commission"]
vals = {r[0]: r for r in ws.iter_rows(values_only=True) if isinstance(r[0], int)}
check("V8 the workbook's commission by grade", all(vals[g][6] == fig["commission_%d" % g] for g in P.GRADES), "")
from docx import Document  # noqa: E402
txt = "\n".join(p.text for p in Document(os.path.join(GD, "low_band_commission_memo.docx")).paragraphs)
check("V9 the memo states the call", ("%d new items" % fig["call"]) in txt or ("{:,} new items".format(fig["call"])) in txt, "")
bad = [r for r in res if not r[1]]
for r in bad:
    print("  FAIL", r[0], r[2])
print("verify: %d of %d" % (len(res) - len(bad), len(res)))
sys.exit(1 if bad else 0)
