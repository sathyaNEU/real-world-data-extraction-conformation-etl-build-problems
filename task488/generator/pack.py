"""The close-out pack: Tessaline's bank export, test events, item responses and item statistics for 2025-26, the
consortium's authoring catalog, framework crosswalk, writing-panel ratings and roster, its scale note, Corvane's
import specification, and two files the answer never reads. tables() returns the shipped tables plus the generator's
truth; write() puts every file into target/."""
from __future__ import annotations

import datetime as dt
import os

import numpy as np
import pandas as pd

import docs as DOCS
import params as P
import world as W

FILES = {
    "bank": "tessaline_item_bank_export_20270331.csv",
    "events": "tessaline_sessions_2025-26.csv",
    "responses": "tessaline_item_responses_2025-26.parquet",
    "stats": "tessaline_item_statistics_2025-26.xlsx",
    "catalog": "qrac_item_authoring_catalog_20270315.xlsx",
    "crosswalk": "qrac_framework_crosswalk_2010_to_2026.csv",
    "panel": "qrac_writing_panel_ratings_2024_2025.csv",
    "roster": "qrac_student_roster_2025-26.csv",
    "scale_note": "qrac_scale_and_cut_scores_2021.pdf",
    "corvane_spec": "corvane_pool_import_specification_v3.pdf",
    "payments": "qrac_workshop_payments_2024_2025.csv",
    "incidents": "tessaline_platform_incidents_2025-26.csv",
}

DISTRICTS = [("QR02", "Ashby Prairie"), ("QR05", "Bellwood"), ("QR07", "Calder Falls"), ("QR11", "Dunmore"),
             ("QR13", "Elk Hollow"), ("QR17", "Fenwick Lake"), ("QR19", "Granite Bluff"), ("QR23", "Harlow"),
             ("QR29", "Ironside"), ("QR31", "Juniper Ridge"), ("QR37", "Kettleby")]
DOMAINS_2010 = {g: (["OA", "NBT", "NF", "MD", "G"] if g <= 5 else (["RP", "NS", "EE", "G", "SP"] if g <= 7 else
                                                                   ["NS", "EE", "F", "G", "SP"])) for g in P.GRADES}
FRAMEWORK = {  # the consortium's 2026 reporting domains
    "lower": [("NSO", "Number sense and operations"), ("ALG", "Algebraic reasoning"), ("FRA", "Fractions"),
              ("MDA", "Measurement and data"), ("GEO", "Geometry")],
    "upper": [("PRR", "Proportional reasoning"), ("NUM", "Number systems"), ("ALG", "Algebraic reasoning"),
              ("GMS", "Geometry and measurement"), ("DAP", "Data and probability")],
}
# 2010 domain to 2026 domain, by grade band; MD in grades 3 to 5 splits by the item's skill tag
MAP_2010 = {"lower": {"OA": "ALG", "NBT": "NSO", "NF": "FRA", "MD": "MDA", "G": "GEO"},
            "upper": {"RP": "PRR", "NS": "NUM", "EE": "ALG", "F": "ALG", "G": "GMS", "SP": "DAP"}}
SPLIT_SKILL = {"lower": ("MD", "measurement word problem", "ALG")}   # 2010 MD items tagged this go to ALG


def band(g):
    return "lower" if g <= 5 else "upper"


def _keys(rng, n, prefix, width, lo, hi):
    vals = rng.choice(np.arange(lo, hi), size=n, replace=False)
    return ["%s%0*d" % (prefix, width, v) for v in vals]


def items_table(Wd, rng):
    """One row per item and grade pool, with the truth columns the checks need kept apart from the shipped ones."""
    rows = []
    keys = iter(_keys(rng, 4000, "TQ", 5, 10000, 99999))
    launch_ids = iter(_keys(rng, 3000, "QM21-", 5, 10000, 99999))
    for g in P.GRADES:
        B = Wd[g]["bank"]
        nk = B["calibrated_qss"].size
        doms = DOMAINS_2010[g]
        wts = np.array([0.24, 0.22, 0.2, 0.19, 0.15])
        for j in range(nk):
            d = doms[int(rng.choice(len(doms), p=wts))]
            rows.append(dict(grade=g, idx=j, source="launch", key=next(keys), auth=next(launch_ids), version=1,
                             dom2010=d, qss=float(B["calibrated_qss"][j]), basis="launch calibration 2021",
                             group="calibrated", truth_b=float(W.theta_of(B["calibrated_qss"][j])),
                             retired=bool(B["retired"][j])))
        nprov = B["prov_true"].size
        for j in range(nprov):
            grp = B["prov_group"][j]
            year = 2024 if rng.random() < 0.46 else 2025
            d = doms[int(rng.choice(len(doms), p=wts))]
            rows.append(dict(grade=g, idx=nk + j, source="workshop %d" % year, key=next(keys),
                             auth="QRAC-M%d-%d-%03d" % (g, year, int(rng.integers(1, 400))), version=1, dom2010=d,
                             qss=float(B["prov_rating_qss"][j]), basis="writing panel rating", group=grp,
                             truth_b=float(B["prov_true"][j]), retired=False))
    T = pd.DataFrame(rows)
    # workshop authoring ids must be unique: redraw the numeric tail where two collide
    seen = set()
    for i in T.index[T.source != "launch"]:
        a = T.at[i, "auth"]
        while a in seen:
            a = a[:-3] + "%03d" % int(rng.integers(1, 999))
        seen.add(a)
        T.at[i, "auth"] = a
    # item types and skill tags (texture; the type does not isolate the panel's misses)
    types = []
    for r in T.itertuples():
        if r.group in ("low_mis", "high_mis"):
            types.append("multi-step word problem")
        elif r.group == "mid_ok":
            types.append("multi-step word problem")
        elif r.group == "right":
            types.append(str(rng.choice(["computation", "one-step word problem"], p=[0.55, 0.45])))
        else:
            types.append(str(rng.choice(["computation", "one-step word problem", "multi-step word problem",
                                         "visual model"], p=[0.38, 0.27, 0.2, 0.15])))
    T["item_type"] = types
    skill = []
    for r in T.itertuples():
        if band(r.grade) == "lower" and r.dom2010 == "MD":
            skill.append(str(rng.choice(["measurement word problem", "data display", "area and perimeter",
                                         "time and money"], p=[0.34, 0.24, 0.24, 0.18])))
        else:
            skill.append({"OA": "operations", "NBT": "place value", "NF": "fractions", "G": "shapes",
                          "RP": "ratios", "NS": "number system", "EE": "expressions", "F": "functions",
                          "SP": "statistics"}.get(r.dom2010, "measurement"))
    T["skill"] = skill
    return T


def ask_metadata(T, rng):
    """Revisions, framework codes, retirements and enemy groups: the import's other contract fields."""
    T = T.copy()
    T["status"] = "operational"
    T["revised"] = False
    T["fw_dom"] = ""
    for i in T.index:
        r = T.loc[i]
        bd = band(int(r.grade))
        base = MAP_2010[bd][r.dom2010]
        if bd == "lower" and r.dom2010 == "MD" and r.skill == SPLIT_SKILL["lower"][1]:
            base = SPLIT_SKILL["lower"][2]
        T.at[i, "xw_dom"] = base
        if r.source == "launch" and not r.retired and rng.random() < 0.095:
            T.at[i, "revised"] = True
            if rng.random() < 0.55:          # re-aligned to another domain in the 2024 review
                others = [c for c, _ in FRAMEWORK[bd] if c != base]
                T.at[i, "fw_dom"] = str(rng.choice(others))
            else:
                T.at[i, "fw_dom"] = base
        elif r.source != "launch":
            T.at[i, "fw_dom"] = base if rng.random() < 0.8 else str(rng.choice([c for c, _ in FRAMEWORK[bd] if c != base]))
    T["true_dom"] = np.where(T.fw_dom != "", T.fw_dom, T.xw_dom)
    T["new_key"] = ""
    cuts = T.grade.map(P.CUT_QSS)
    can = T.revised & (T.qss > cuts + 8) & (T.source == "launch")
    pick = can & (rng.random(len(T)) < 0.62)
    pool = iter(_keys(rng, int(pick.sum()) + 50, "TQ", 5, 10000, 99999))
    used = set(T.key)
    for i in T.index[pick]:
        k = next(pool)
        while k in used:
            k = next(pool)
        used.add(k)
        T.at[i, "new_key"] = k
    # retired launch items (2024): in both systems, never administered in 2025-26, not imported
    T.loc[T.retired.astype(bool), "status"] = "retired"
    # enemy groups at the authoring-version level
    groups = []
    gid = 0
    for g in P.GRADES:
        pool = T[(T.grade == g)].index.to_numpy()
        n_groups = int(rng.integers(9, 15))
        revised_pool = T[(T.grade == g) & T.revised].index.to_numpy()
        for _ in range(n_groups):
            size = int(rng.choice([2, 2, 2, 3]))
            mem = list(rng.choice(pool, size=size, replace=False))
            if len(revised_pool) and rng.random() < 0.55:     # revisions often came out of an enemy review
                mem[0] = int(rng.choice(revised_pool))
            if rng.random() < 0.18:      # a cross-grade member, from the next grade up or down
                og = g + (1 if g < 8 else -1)
                mem.append(int(rng.choice(T[T.grade == og].index.to_numpy())))
            gid += 1
            for m in mem:
                groups.append((gid, int(m)))
    G = pd.DataFrame(groups, columns=["gid", "row"]).drop_duplicates()
    return T, G


def group_of(G):
    """Each item's enemy group: the first group it was put in (an item sits in one group)."""
    mem = {}
    for gid, sub in G.sort_values(["gid", "row"]).groupby("gid"):
        for row in sub.row:
            mem.setdefault(int(row), "EG-%03d" % (100 + gid))
    return mem


def versions(T, G):
    """The authoring catalog's version history, drawn on its own stream so the delivery data never moves: the 2024
    alignment and accessibility review made version 2 of every revised launch item, a 2025 accessibility pass made a
    version 3 of about half of them, and a group set on version 1 was cleared by the 2024 review for some items."""
    rng = np.random.default_rng(P.SEED_PACK + 23)
    T = T.copy()
    mem = group_of(G)
    T["eg"] = [mem.get(int(i), "") for i in T.index]
    u = rng.random((len(T), 2))
    T["eg_on_v1"] = T.revised & (T.eg != "") & (u[:, 0] < 0.8)
    T["eg_removed"] = T.eg_on_v1 & (u[:, 1] < 0.25)
    # the 2025 accessibility pass: a version 3 of a revised launch item, a version 2 of a 2024 workshop item, drawn
    # until it moves every domain count of every grade for a reader who takes the active row as the whole record
    plain = np.array([MAP_2010[band(int(r.grade))][r.dom2010] for r in T.itertuples()])
    truth = T.true_dom.to_numpy()
    cand = (T.revised | (T.source == "workshop 2024")).to_numpy()
    acc = np.zeros(len(T), dtype=bool)
    for g in P.GRADES:
        idx = np.where(cand & (T.grade == g).to_numpy())[0]
        doms = [c for c, _ in FRAMEWORK[band(g)]]
        for _ in range(2000):
            sel = idx[rng.random(len(idx)) < 0.5]
            mv = sel[plain[sel] != truth[sel]]
            if all((plain[mv] == d).sum() != (truth[mv] == d).sum() for d in doms):
                break
        acc[sel] = True
    T["acc25"] = acc
    return T


def ask_truth(T, G):
    """Ask A: imported items per grade and 2026 domain, one per item (a re-keyed revision once, under its current
    key). Ask B: enemy sets per grade and the items they hold."""
    imp = T[T.status == "operational"]
    dom = {(g, d): int(((imp.grade == g) & (imp.true_dom == d)).sum()) for g in P.GRADES
           for d, _ in FRAMEWORK[band(g)]}
    return dom, enemy_truth(T, G)


def enemy_truth(T, G):
    """Corvane's sets: a group's imported (operational) items in one grade's pool, two or more to a set; a membership
    the 2024 review cleared is gone."""
    out = {}
    G = G.sort_values(["gid", "row"]).drop_duplicates("row", keep="first")     # an item sits in one group, its first
    G = G[~T.loc[G.row, "eg_removed"].to_numpy()]
    for g in P.GRADES:
        n_sets = n_items = 0
        for gid, sub in G.groupby("gid"):
            rows = T.loc[sub.row]
            keep = rows[(rows.grade == g) & (rows.status == "operational")]
            if len(keep) >= 2:
                n_sets += 1
                n_items += len(keep)
        out[g] = (n_sets, n_items)
    return out


CLUSTERS = {"OA": "ABCD", "NBT": "AB", "NF": "ABC", "MD": "ABCD", "G": "A", "RP": "A", "NS": "ABC", "EE": "ABC",
            "F": "AB", "SP": "AB"}


def std_code(g, dom, rng):
    cl = CLUSTERS[dom]
    return "%d.%s.%s.%d" % (g, dom, cl[int(rng.integers(0, len(cl)))], int(rng.integers(1, 5)))


def central_to_utc(d, hh, mm):
    # Central time: daylight until 2 Nov 2025 and from 8 Mar 2026, standard between
    off = 6 if dt.date(2025, 11, 2) <= d < dt.date(2026, 3, 8) else 5
    return dt.datetime(d.year, d.month, d.day, hh, mm) + dt.timedelta(hours=off)


def school_days(d0, d1):
    out, d = [], d0
    while d <= d1:
        if d.weekday() < 5:
            out.append(d)
        d += dt.timedelta(days=1)
    return out


def tables():
    Wd = W.simulate()
    rng = np.random.default_rng(P.SEED_PACK)
    T = items_table(Wd, rng)
    T, G = ask_metadata(T, rng)
    T = versions(T, G)
    T["std2010"] = [std_code(int(r.grade), r.dom2010, rng) for r in T.itertuples()]
    # ---- events and responses
    ev_rows, resp_parts = [], []
    stud_rows = []
    n_stud = sum(P.STUDENTS.values())
    n_ev = sum(len(Wd[g]["run"]["events"]) for g in P.GRADES)
    skey_all = iter(_keys(rng, 2 * n_stud, "S", 8, 10_000_000, 99_999_999))
    cid_all = iter(_keys(rng, n_stud, "", 10, 1_000_000_000, 9_999_999_999))
    evid_all = iter(_keys(rng, n_ev, "E", 9, 100_000_000, 999_999_999))
    for g in P.GRADES:
        R = Wd[g]["run"]
        ns = P.STUDENTS[g]
        dist = rng.choice(len(DISTRICTS), size=ns, p=np.array([9, 7, 12, 8, 10, 6, 11, 9, 8, 13, 7]) / 100.0)
        school = ["%s-%02d" % (DISTRICTS[d][0], int(rng.integers(1, 6))) for d in dist]
        skey = [next(skey_all) for _ in range(ns)]
        cid = [next(cid_all) for _ in range(ns)]
        mover = rng.random(ns) < 0.018
        new_dist = np.where(mover, (dist + rng.integers(1, len(DISTRICTS), ns)) % len(DISTRICTS), dist)
        skey2 = [next(skey_all) for _ in range(ns)]
        for s in range(ns):
            stud_rows.append(dict(student_key=skey[s], consortium_id=cid[s], district_code=DISTRICTS[dist[s]][0],
                                  school_code=school[s], grade=g, entry_date="2025-08-27",
                                  exit_date="2026-01-09" if mover[s] else ""))
            if mover[s]:
                stud_rows.append(dict(student_key=skey2[s], consortium_id=cid[s],
                                      district_code=DISTRICTS[new_dist[s]][0],
                                      school_code="%s-%02d" % (DISTRICTS[new_dist[s]][0], int(rng.integers(1, 6))),
                                      grade=g, entry_date="2026-01-12", exit_date=""))
        evid = [next(evid_all) for _ in range(len(R["events"]))]
        rep_eap = np.array([e[3] for e in R["events"]])
        days = [school_days(w[1], w[2]) for w in P.WINDOWS]
        for k, e in enumerate(R["events"]):
            s, w = e[0], e[1]
            wname = P.WINDOWS[w][0]
            day = days[w][int(rng.integers(0, len(days[w])))]
            st = central_to_utc(day, int(rng.integers(8, 14)), int(rng.choice([0, 5, 10, 15, 20, 30, 40, 45, 50])))
            dur = dt.timedelta(minutes=int(18 + e[5] * 1.1 + rng.integers(0, 25)))
            key = skey2[s] if (mover[s] and w > 0) else skey[s]
            dcode = DISTRICTS[new_dist[s]][0] if (mover[s] and w > 0) else DISTRICTS[dist[s]][0]
            ss = int(np.clip(round(P.SCALE_A + P.SCALE_B * e[3]), P.LOSS, P.HOSS))
            sem = int(round(P.SCALE_B * e[4]))
            status = e[6]
            ev_rows.append(dict(event_id=evid[k], student_key=key, grade=g, district_code=dcode,
                                window=wname, started_utc=st.strftime("%Y-%m-%dT%H:%M:%SZ"),
                                finished_utc=(st + dur).strftime("%Y-%m-%dT%H:%M:%SZ"), status=status,
                                items_answered=e[5], scale_score="" if status == "invalidated" else ss,
                                sem="" if status == "invalidated" else sem, _rep=rep_eap[k], _true=e[2]))
        nk = R["nk"]
        gi = T[T.grade == g].set_index("idx")
        ikey = gi.key.reindex(np.arange(len(gi))).to_numpy()
        ev_ids = np.array(evid)
        n = R["resp_item"].size
        secs = np.clip(np.round(np.exp(rng.normal(3.75, 0.55, n))), 4, 600).astype(np.int16)
        nkey = gi.new_key.reindex(np.arange(len(gi))).to_numpy()
        ev_win = np.array([e[1] for e in R["events"]])
        use_new = (nkey[R["resp_item"]] != "") & (ev_win[R["resp_event"]] > 0)
        keys_out = np.where(use_new, nkey[R["resp_item"]], ikey[R["resp_item"]])
        df = pd.DataFrame({"event_id": ev_ids[R["resp_event"]], "sequence": R["resp_seq"].astype(np.int8),
                           "item_key": keys_out, "score": R["resp_score"].astype("float64"),
                           "response_seconds": secs})
        # a resumed session re-delivers the item that was on screen: first row unscored, at the timeout
        n_ev = len(R["events"])
        res_ev = rng.choice(n_ev, size=int(P.RESUMED_SHARE * n_ev), replace=False)
        lens = np.array([e[5] for e in R["events"]])
        dup = []
        first_pos = np.searchsorted(R["resp_event"], res_ev)
        for e_i, fp in zip(res_ev, first_pos):
            sq = int(rng.integers(2, max(3, lens[e_i])))
            r = fp + sq - 1
            if r < n and R["resp_event"][r] == e_i:
                dup.append(dict(event_id=ev_ids[e_i], sequence=np.int8(sq), item_key=keys_out[r],
                                score=np.nan, response_seconds=np.int16(1200)))
        df = pd.concat([pd.DataFrame(dup), df], ignore_index=True)
        resp_parts.append(df)
    E = pd.DataFrame(ev_rows)
    Rsp = pd.concat(resp_parts, ignore_index=True)
    Rsp["score"] = Rsp["score"].astype("Int8")
    Rsp = Rsp.sort_values(["event_id", "sequence", "score"], na_position="first", kind="mergesort").reset_index(drop=True)
    E = E.sort_values(["grade", "window", "event_id"], kind="mergesort").reset_index(drop=True)
    roster = pd.DataFrame(stud_rows).sort_values(["district_code", "student_key"], kind="mergesort").reset_index(drop=True)
    return dict(T=T, G=G, E=E, R=Rsp, roster=roster, Wd=Wd, rng=rng)


INTERACTION = ["MC4", "MC4", "MC4", "NUM", "NUM", "MS", "EQN", "DND"]


def catalog_rows(T, G, rng):
    """The consortium's authoring catalog: one row per authored version. Version 1 is the item's full record; a later
    version records the fields it changed, listed in changed_fields (a listed field left blank was cleared), so an
    enemy group set on version 1 and a framework code set on version 2 stand on the versions after them."""
    rows = []
    fw_n = {}
    for i, r in T.iterrows():
        g = int(r.grade)
        eg = r.eg
        code = ""
        if r.fw_dom:
            fw_n[(r.fw_dom, g)] = fw_n.get((r.fw_dom, g), 0) + 1
            code = "%s.%d.%d" % (r.fw_dom, g, 1 + (fw_n[(r.fw_dom, g)] * 7) % 9)
        ident = dict(authoring_id=r.auth, item_key=r.key, grade=g, source=r.source)
        if r.revised:
            rows.append(dict(ident, version=1, status="superseded", replaces="", changed_fields="",
                             last_reviewed="2021-03-18", item_type=r.item_type, skill_tag=r.skill,
                             framework_code_2026="", enemy_group=eg if r.eg_on_v1 else ""))
            look = sorted(rng.choice(["art", "alt_text", "layout"], size=int(rng.integers(1, 3)), replace=False))
            ch = list(look) + ["framework_code_2026"]
            if (eg and not r.eg_on_v1) or r.eg_removed:
                ch.append("enemy_group")
            key2 = r.new_key or r.key
            rows.append(dict(ident, item_key=key2, version=2, status="superseded" if r.acc25 else "active",
                             replaces="%s v1" % r.auth, changed_fields="; ".join(ch),
                             last_reviewed=str(dt.date(2024, 6, 3) + dt.timedelta(days=int(rng.integers(0, 60)))),
                             item_type="", skill_tag="", framework_code_2026=code,
                             enemy_group=eg if (eg and not r.eg_on_v1) else ""))
            if r.acc25:
                rows.append(access_row(ident, key2, 3, r.auth, rng))
        else:
            st = "retired" if r.status == "retired" else ("superseded" if r.acc25 else "active")
            rev = "2021-03-18" if r.source == "launch" else ("2024-08-%02d" % int(rng.integers(5, 28)) if
                                                              r.source.endswith("2024") else
                                                              "2025-08-%02d" % int(rng.integers(4, 27)))
            rows.append(dict(ident, version=1, status=st, replaces="", changed_fields="", last_reviewed=rev,
                             item_type=r.item_type, skill_tag=r.skill, framework_code_2026=code, enemy_group=eg))
            if r.acc25:
                rows.append(access_row(ident, r.key, 2, r.auth, rng))
    C = pd.DataFrame(rows)
    C = C[["authoring_id", "version", "item_key", "grade", "source", "status", "replaces", "changed_fields",
           "last_reviewed", "item_type", "skill_tag", "framework_code_2026", "enemy_group"]]
    return C.sort_values(["grade", "authoring_id", "version"], kind="mergesort").reset_index(drop=True)


def access_row(ident, key, version, auth, rng):
    """A version from the 2025 accessibility pass: art, alt text or text-to-speech markup, nothing else."""
    return dict(ident, item_key=key, version=version, status="active", replaces="%s v%d" % (auth, version - 1),
                changed_fields=str(rng.choice(["tts_markup", "alt_text", "art", "tts_markup; alt_text"])),
                last_reviewed=str(dt.date(2025, 5, 12) + dt.timedelta(days=int(rng.integers(0, 50)))),
                item_type="", skill_tag="", framework_code_2026="", enemy_group="")


def crosswalk_rows(T):
    rows = []
    for code in sorted(set(T.std2010)):
        g = int(code.split(".")[0])
        dom = code.split(".")[1]
        bd = band(g)
        target = MAP_2010[bd][dom]
        if bd == "lower" and dom == "MD":
            rows.append(dict(standard_code_2010=code, framework_domain_2026=SPLIT_SKILL["lower"][2],
                             condition="skill tag: measurement word problem"))
            rows.append(dict(standard_code_2010=code, framework_domain_2026=target, condition="other skill tags"))
        else:
            rows.append(dict(standard_code_2010=code, framework_domain_2026=target, condition=""))
    return pd.DataFrame(rows)


def panel_rows(T, rng):
    rows = []
    for r in T[T.source != "launch"].itertuples():
        p = int(round(r.qss))
        d1, d2 = int(rng.integers(-6, 7)), int(rng.integers(-6, 7))
        d3 = -d1 - d2 + int(rng.integers(-1, 2))
        vals = [p + d1, p + d2, p + d3]
        if int(round(np.mean(vals))) != p:
            vals[2] = 3 * p - vals[0] - vals[1]
        when = "2024-07-%02d" % int(rng.integers(15, 26)) if r.source.endswith("2024") else "2025-07-%02d" % int(rng.integers(14, 25))
        panel = rng.choice(["panel A", "panel B", "panel C"])
        for k, v in enumerate(vals):
            rows.append(dict(authoring_id=r.auth, workshop=r.source, panel=str(panel), rater=k + 1, rating_qss=v,
                             rated_on=when))
    return pd.DataFrame(rows)


def stats_frames(tb):
    T, E, R = tb["T"], tb["E"], tb["R"]
    r = R[R.score.notna()].merge(E[["event_id", "window", "status"]], on="event_id")
    r = r[r.status != "invalidated"]
    out = {}
    for w in [x[0] for x in P.WINDOWS]:
        s = r[r.window == w].groupby("item_key").agg(administrations=("score", "size"),
                                                     correct=("score", "sum"), median_seconds=("response_seconds", "median"))
        s["percent_correct"] = (100.0 * s.correct / s.administrations).round(1)
        s = s.join(T.set_index("key")[["grade"]]).reset_index().rename(columns={"index": "item_key"})
        s = s[["item_key", "grade", "administrations", "percent_correct", "median_seconds"]]
        s = s.sort_values(["grade", "item_key"], kind="mergesort").reset_index(drop=True)
        out[w] = s.rename(columns={"item_key": "Item", "grade": "Grade", "administrations": "Administrations",
                                   "percent_correct": "Percent correct", "median_seconds": "Median seconds"})
    return out


# the workshop's stipend schedule by year: a day rate, a rate per accepted item, and mileage to the consortium office
# in Dunmore at the federal business rate of the year (cents a mile), paid on the next two payroll runs or on invoice
DAY_RATE = {"workshop 2024": 145, "workshop 2025": 155}
ITEM_RATE = {"workshop 2024": 80, "workshop 2025": 85}
MILE_CENTS = {"workshop 2024": 67, "workshop 2025": 70}
PAYROLL = {"workshop 2024": ("2024-08-09", "2024-08-23"), "workshop 2025": ("2025-08-08", "2025-08-22")}
MILES_ONE_WAY = {"QR02": 41, "QR05": 17, "QR07": 58, "QR11": 0, "QR13": 33, "QR17": 26, "QR19": 64, "QR23": 9,
                 "QR29": 47, "QR31": 22, "QR37": 36}


def distractor_frames(T, rng):
    pay = []
    tx = np.random.default_rng(P.SEED_PACK + 19)
    for wk in ("workshop 2024", "workshop 2025"):
        for g in P.GRADES:
            n_w = int(rng.integers(4, 8))
            acc = int(((T.source == wk) & (T.grade == g)).sum())
            split = rng.multinomial(acc, np.ones(n_w) / n_w)
            for k in range(n_w):
                sub = int(split[k] + rng.integers(1, 6))
                home = DISTRICTS[int(tx.integers(0, len(DISTRICTS)))][0]
                days = int(tx.choice([4, 3], p=[0.88, 0.12]))
                claims = MILES_ONE_WAY[home] > 0 and tx.random() < 0.8
                miles = 2 * MILES_ONE_WAY[home] * days if claims else 0
                mileage = round(miles * MILE_CENTS[wk] / 100, 2)
                retired = tx.random() < 0.12
                if retired:
                    y = int(wk[-4:])
                    paid = (dt.date(y, 8, 12) + dt.timedelta(days=int(tx.integers(0, 32)))).isoformat()
                else:
                    paid = PAYROLL[wk][int(tx.random() < 0.3)]
                total = days * DAY_RATE[wk] + int(split[k]) * ITEM_RATE[wk] + mileage
                pay.append(dict(writer_code="W%s%d%02d" % (wk[-2:], g, k + 1), workshop=wk, grade=g,
                                home_district=home, days_attended=days, items_submitted=sub,
                                items_accepted=int(split[k]), day_rate_usd=DAY_RATE[wk],
                                item_rate_usd=ITEM_RATE[wk], miles_claimed=miles, mileage_usd="%.2f" % mileage,
                                total_paid_usd="%.2f" % total, paid_on=paid,
                                paid_via="consultant invoice" if retired else "district payroll"))
    inc = []
    base = dt.datetime(2025, 9, 16, 14, 0)
    for k in range(int(rng.integers(9, 14))):
        st = base + dt.timedelta(days=int(rng.integers(0, 240)), minutes=int(rng.integers(0, 420)))
        dists = sorted(set(str(d[0]) for d in [DISTRICTS[int(i)] for i in rng.choice(len(DISTRICTS), int(rng.integers(1, 4)), replace=False)]))
        inc.append(dict(incident_id="INC-%05d" % int(rng.integers(20000, 99999)),
                        started_utc=st.strftime("%Y-%m-%dT%H:%M:%SZ"),
                        resolved_utc=(st + dt.timedelta(minutes=int(rng.integers(12, 160)))).strftime("%Y-%m-%dT%H:%M:%SZ"),
                        districts=";".join(dists), sessions_interrupted=int(rng.integers(8, 260)),
                        summary=str(rng.choice(["session timeouts on the delivery service",
                                                "slow item rendering in the secure browser",
                                                "login service degraded", "scoring queue delayed"]))))
    return pd.DataFrame(pay), pd.DataFrame(inc).sort_values("started_utc").reset_index(drop=True)


def shipped(tb):
    """The tables exactly as they ship (no truth columns)."""
    T, E, R = tb["T"], tb["E"], tb["R"]
    rng = np.random.default_rng(P.SEED_PACK + 7)
    inter = [INTERACTION[int(i)] for i in rng.integers(0, len(INTERACTION), len(T))]
    last = []
    for r in T.itertuples():
        if r.status == "retired":
            last.append("2024-05-%02d" % int(rng.integers(6, 24)))
        else:
            last.append("2026-05-%02d" % int(rng.integers(11, 23)))
    bank = pd.DataFrame({"item_key": T.key, "grade": T.grade, "status": T.status, "standard_code": T.std2010,
                         "interaction_type": inter,
                         "difficulty_qss": T.qss.round(1), "difficulty_basis": T.basis,
                         "first_operational": np.where(T.source == "launch", "2021 fall",
                                                       np.where(T.source == "workshop 2024", "2024 fall", "2025 fall")),
                         "last_administered": last})
    rk = T[T.new_key != ""]
    extra = bank[bank.item_key.isin(set(rk.key))].copy()
    extra["item_key"] = extra.item_key.map(dict(zip(rk.key, rk.new_key)))
    extra["first_operational"] = "2024 fall"
    bank = pd.concat([bank, extra], ignore_index=True)
    bank = bank.sort_values(["grade", "item_key"], kind="mergesort").reset_index(drop=True)
    events = E[["event_id", "student_key", "grade", "district_code", "window", "started_utc", "finished_utc", "status",
                "items_answered", "scale_score", "sem"]].copy()
    resp = R[["event_id", "sequence", "item_key", "score", "response_seconds"]].copy()
    cat = catalog_rows(T, tb["G"], np.random.default_rng(P.SEED_PACK + 11))
    xw = crosswalk_rows(T)
    panel = panel_rows(T, np.random.default_rng(P.SEED_PACK + 13))
    pay, inc = distractor_frames(T, np.random.default_rng(P.SEED_PACK + 17))
    stats = stats_frames(tb)
    return dict(bank=bank, events=events, responses=resp, catalog=cat, crosswalk=xw, panel=panel,
                roster=tb["roster"], payments=pay, incidents=inc, stats=stats)


def _xlsx(path, sheets, when, company, author):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    first = True
    for name, df in sheets.items():
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        ws.append(list(df.columns))
        for c in ws[1]:
            c.font = Font(bold=True)
        for row in df.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in row])
        ws.freeze_panes = "A2"
        for k, col in enumerate(df.columns):
            ws.column_dimensions[openpyxl.utils.get_column_letter(k + 1)].width = max(10, min(28, len(col) + 3))
    wb.save(path)
    DOCS.fixed_xlsx(path, when=when, company=company, author=author)


def write(target, s):
    os.makedirs(target, exist_ok=True)
    n = {}

    def csv(key, df):
        df.to_csv(os.path.join(target, FILES[key]), index=False, lineterminator="\n")
        n[FILES[key]] = len(df)
    csv("bank", s["bank"])
    csv("events", s["events"])
    csv("crosswalk", s["crosswalk"])
    csv("panel", s["panel"])
    csv("roster", s["roster"])
    csv("payments", s["payments"])
    csv("incidents", s["incidents"])
    import pyarrow as pa
    import pyarrow.parquet as pq
    tbl = pa.Table.from_pandas(s["responses"], preserve_index=False)
    pq.write_table(tbl, os.path.join(target, FILES["responses"]), compression="zstd", use_dictionary=True)
    n[FILES["responses"]] = len(s["responses"])
    _xlsx(os.path.join(target, FILES["stats"]), {w: s["stats"][w] for w in [x[0] for x in P.WINDOWS]},
          (2026, 6, 4, 13, 20, 0), b"Tessaline Learning Systems", b"Tessaline Reporting")
    n[FILES["stats"]] = sum(len(v) for v in s["stats"].values())
    _xlsx(os.path.join(target, FILES["catalog"]), {"items": s["catalog"]}, (2027, 3, 15, 15, 5, 0),
          b"Quarrystone Regional Assessment Consortium", b"QRAC Assessment Office")
    n[FILES["catalog"]] = len(s["catalog"])
    DOCS.scale_note(os.path.join(target, FILES["scale_note"]))
    DOCS.corvane_spec(os.path.join(target, FILES["corvane_spec"]))
    return n
