"""The pack's documents and lookup files: capacity plan, planning basis, change log, fleet register, route
master, close-out summary, data dictionary, the two distractors and the export manifest."""
import csv
import datetime as dt
import os
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import world as W
import sim
from records import BLOCKS, route_list
from fileio import fix_pdf, fix_workbook_props, normalise_zip

STATION_NAME = {
    "LIN": "Milano Linate", "TRN": "Torino", "VRN": "Verona", "VCE": "Venezia", "BLQ": "Bologna",
    "CTA": "Catania", "PMO": "Palermo", "NAP": "Napoli", "BRI": "Bari", "BDS": "Brindisi",
    "SUF": "Lamezia Terme", "REG": "Reggio Calabria", "CAG": "Cagliari", "TPS": "Trapani", "AHO": "Alghero",
    "CRV": "Crotone", "PSR": "Pescara", "CIY": "Comiso", "OLB": "Olbia",
}
REGION = {"CTA": "Sicilia", "PMO": "Sicilia", "TPS": "Sicilia", "CIY": "Sicilia", "NAP": "Campania",
          "BRI": "Puglia", "BDS": "Puglia", "SUF": "Calabria", "REG": "Calabria", "CRV": "Calabria",
          "CAG": "Sardegna", "AHO": "Sardegna", "OLB": "Sardegna", "PSR": "Abruzzo"}
TYPE_LONG = {"319": "A319", "320": "A320", "32N": "A320neo", "321": "A321", "32Q": "A321neo"}

_SS = getSampleStyleSheet()
BODY = ParagraphStyle("body", parent=_SS["Normal"], fontName="Helvetica", fontSize=9.2, leading=12.2, spaceAfter=5)
SMALL = ParagraphStyle("small", parent=BODY, fontSize=7.6, leading=9.6, spaceAfter=3)
H1 = ParagraphStyle("h1", parent=BODY, fontName="Helvetica-Bold", fontSize=12.5, leading=15, spaceAfter=6)
H2 = ParagraphStyle("h2", parent=BODY, fontName="Helvetica-Bold", fontSize=9.6, leading=12.5, spaceBefore=6,
                    spaceAfter=3)


def _pdf(path, story, title, author, pagesize=A4, when=dt.datetime(2025, 11, 28, 9, 0)):
    def meta(canvas, doc):
        canvas.setTitle(title)
        canvas.setAuthor(author)
        canvas.setCreator(author)
        canvas.setProducer("Navigli Air document services")
        canvas.setSubject(title)
        canvas.setFont("Helvetica", 7)
        canvas.drawRightString(doc.pagesize[0] - 15 * mm, 9 * mm, "page %d" % doc.page)
    doc = SimpleDocTemplate(path, pagesize=pagesize, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=15 * mm,
                            bottomMargin=15 * mm, invariant=1, title=title, author=author, creator=author)
    doc.build(story, onFirstPage=meta, onLaterPages=meta)
    fix_pdf(path, when)


def _table(data, widths, header_rows=1, font=7.6, align_right_from=1):
    t = Table(data, colWidths=widths, repeatRows=header_rows)
    st = [("FONT", (0, 0), (-1, -1), "Helvetica", font), ("FONT", (0, 0), (-1, header_rows - 1), "Helvetica-Bold", font),
          ("LINEBELOW", (0, header_rows - 1), (-1, header_rows - 1), 0.6, colors.black),
          ("ALIGN", (align_right_from, header_rows), (-1, -1), "RIGHT"),
          ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6),
          ("TOPPADDING", (0, 0), (-1, -1), 1.6)]
    t.setStyle(TableStyle(st))
    return t


def season_dates(season):
    return sim.season_dates(season)


# ------------------------------------------------------------------ route-level facts used by several documents
def route_rows(S):
    """One dict per route with its gauge, frequency and season seats in S24, S25 and S26."""
    s26 = season_dates("S26")
    out = []
    for b, st, kind in route_list():
        if kind == "target":
            spec = [t for t in W.TARGETS if t[0] == st][0]
            freq = spec[1]
            t24 = t25 = "319"
            sw = W.TARGET_SWITCH.get(st)
            t26 = "32Q"
        elif kind == "past":
            spec = [p for p in W.PAST if p[0] == b and p[1] == st][0]
            freq = spec[2]
            t24, t25, t26, sw = spec[3], spec[4], spec[4], None
        else:
            spec = [p for p in W.REF if p[0] == b and p[1] == st][0]
            freq = spec[2]
            t24 = t25 = t26 = spec[3]
            sw = None
        seats25 = W.SEATS[t25] * freq * 2 * 210
        if sw is not None:
            before = sum(1 for d in s26 if d < sw)
            seats26 = (W.SEATS["319"] * before + W.SEATS[t26] * (210 - before)) * freq * 2
        else:
            seats26 = W.SEATS[t26] * freq * 2 * 210
        out.append(dict(base=b, station=st, kind=kind, freq=freq, t24=t24, t25=t25, t26=t26, switch=sw,
                        seats24=W.SEATS[t24] * freq * 2 * 210, seats25=seats25, seats26=seats26,
                        dist=W.DISTANCE_KM[(b, st)]))
    return out


# ------------------------------------------------------------------ capacity plan, summer 2026, version 3
def capacity_plan(S, path):
    rr = route_rows(S)
    wb = Workbook()
    ws = wb.active
    ws.title = "Cover"
    lines = [
        ("Navigli Air - Network Planning", True),
        ("Summer 2026 segment capacity plan", True),
        ("Version 3, 21/11/2025", False),
        ("Owner: Ugo Pacomio, Head of Network Planning", False),
        ("Scope: scheduled departures of the IATA summer season 2026 (29/03/2026 - 24/10/2026) against summer "
         "2025 (30/03/2025 - 25/10/2025) as flown", False),
        ("Seats are the physical seats of the planned aircraft type; departures are scheduled departures, both "
         "directions of each route.", False),
        ("Daily pattern: every route keeps its summer 2025 departure times and operates every day of the season.",
         False),
    ]
    for k, (txt, bold) in enumerate(lines, start=1):
        ws.cell(row=k, column=1, value=txt).font = Font(bold=bold, size=12 if k <= 2 else 10)
    ws.column_dimensions["A"].width = 110

    ws = wb.create_sheet("Routes")
    head = ["Route", "Base", "Station", "Type S25", "Seats/dep S25", "Dep/day each way S25", "Season seats S25",
            "Type S26", "Seats/dep S26", "S26 type from", "Dep/day each way S26", "Season seats S26",
            "Change in seats"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE4EE")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rr:
        frm = "01/06/2026 (A319 to 31/05/2026)" if r["switch"] else ("29/03/2026" if r["t26"] != r["t25"] else "")
        ws.append(["%s-%s" % (r["base"], r["station"]), r["base"], r["station"], TYPE_LONG[r["t25"]],
                   W.SEATS[r["t25"]], r["freq"], r["seats25"], TYPE_LONG[r["t26"]], W.SEATS[r["t26"]], frm,
                   r["freq"], r["seats26"], r["seats26"] - r["seats25"]])
    n = len(rr) + 1
    ws.append([])
    ws.append(["Total", "", "", "", "", "", "=SUM(G2:G%d)" % n, "", "", "", "", "=SUM(L2:L%d)" % n,
               "=SUM(M2:M%d)" % n])
    for col, w in zip("ABCDEFGHIJKLM", [10, 6, 8, 9, 9, 10, 12, 9, 9, 26, 10, 12, 11]):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet("Gauge changes")
    ws.append(["Route", "From", "To", "Effective", "Seats added in season", "Aircraft"])
    for c in ws[1]:
        c.font = Font(bold=True)
    alloc = {"CTA": "I-NQAA, I-NQAB", "PMO": "I-NQAC, I-NQAD", "NAP": "I-NQAE", "BRI": "I-NQAF", "CAG": "I-NQAG",
             "REG": "I-NQAH (with TPS)", "TPS": "I-NQAH (with REG)", "BDS": "I-NQAL", "SUF": "I-NQAM"}
    total = 0
    for r in rr:
        if r["kind"] != "target":
            continue
        eff = "01/06/2026" if r["switch"] else "29/03/2026"
        ws.append(["LIN-%s" % r["station"], "A319 (144)", "A321neo (232)", eff, r["seats26"] - r["seats25"],
                   alloc[r["station"]]])
        total += r["seats26"] - r["seats25"]
    ws.append(["Nine Linate routes", "", "", "", total, ""])
    ws.append([])
    ws.append(["A321neo deliveries: two in January, four in February, two in March and the last two in May 2026 "
               "(the two April slots moved to May after the engine supplier's notice of 6 November 2025)."])
    for col, w in zip("ABCDEF", [18, 12, 14, 12, 18, 28]):
        ws.column_dimensions[col].width = w

    ws = wb.create_sheet("Notes")
    for row in [["Version", "Date", "Change"],
                ["1", "22/09/2025", "first plan for summer 2026"],
                ["2", "24/10/2025", "A321neo on the nine Linate routes from the season start"],
                ["3", "21/11/2025", "Brindisi and Lamezia Terme on the A321neo from 01/06/2026 (last two "
                                    "deliveries in May); A319 on both until 31/05/2026"]]:
        ws.append(row)
    ws["A1"].font = ws["B1"].font = ws["C1"].font = Font(bold=True)
    ws.column_dimensions["C"].width = 90
    fix_workbook_props(wb, "Summer 2026 segment capacity plan v3", "Ugo Pacomio", dt.datetime(2025, 11, 21, 17, 4))
    wb.save(path)
    normalise_zip(path)
    return total


# ------------------------------------------------------------------ the planning basis (the estimand's home)
def planning_basis(path):
    st = []
    st.append(Paragraph("Navigli Air - Network Planning", SMALL))
    st.append(Paragraph("Summer 2026 planning basis: the nine Linate gauge changes", H1))
    st.append(Paragraph("Document NP-PB-26S, issue 2, 28 November 2025. Owner: Ugo Pacomio, Head of Network "
                        "Planning. Distribution: Commercial, Revenue Management, Ground Operations, Finance.", SMALL))
    st.append(Spacer(1, 4 * mm))
    paras = [
        ("1. Use", "The summer 2026 budget and the outstation handling tender are built on this basis for the nine "
                   "Linate routes that change gauge for summer 2026. Ground Operations (Vincenza Passalacqua) "
                   "issues the tender on 15 January 2026."),
        ("2. Routes", "Milano Linate to and from Catania (CTA), Palermo (PMO), Napoli (NAP), Bari (BRI), Brindisi "
                      "(BDS), Lamezia Terme (SUF), Reggio Calabria (REG), Cagliari (CAG) and Trapani (TPS). A route "
                      "counts both directions."),
        ("3. Seasons", "Summer 2025 is the IATA summer season 30 March to 25 October 2025, as flown. Summer 2026 is "
                       "the IATA summer season 29 March to 24 October 2026, as scheduled in capacity plan version 3. "
                       "Both seasons have 210 days, and summer 2026 repeats summer 2025's departures, the same "
                       "flights on the same weekdays at the same times."),
        ("4. Passengers carried", "Passengers carried on a route are the passengers flown on its departures as "
                                  "closed out in departure control. Lap infants are recorded separately and are not "
                                  "passengers carried; no-shows are not passengers carried."),
        ("5. Capacity", "Summer 2026 capacity is capacity plan version 3. The nine routes move from the A319 (144 "
                        "seats) to the A321neo (232 seats) from the season start, except Brindisi and Lamezia Terme, "
                        "which keep the A319 until 31 May 2026 and fly the A321neo from 1 June 2026. Seats are sold "
                        "up to the physical seats; we do not overbook domestic flights."),
        ("6. Market", "No market growth is planned for summer 2026: the passengers who wanted to fly the nine routes "
                      "in summer 2025, on the departures they wanted, are the passengers summer 2026 plans for."),
        ("7. The figure", "Passenger revenue in the budget and the passenger volumes in the handling tender are "
                          "built on the additional passengers carried on the nine routes in summer 2026 over summer "
                          "2025."),
    ]
    for head, txt in paras:
        st.append(Paragraph(head, H2))
        st.append(Paragraph(txt, BODY))
    st.append(Spacer(1, 4 * mm))
    st.append(_table([["Issue", "Date", "Change"], ["1", "07/11/2025", "first issue"],
                      ["2", "28/11/2025", "section 5 aligned to capacity plan version 3 (Brindisi and Lamezia Terme "
                                          "from 1 June 2026)"]],
                     [14 * mm, 22 * mm, 140 * mm], align_right_from=9))
    _pdf(path, st, "Summer 2026 planning basis NP-PB-26S issue 2", "Ugo Pacomio")


# ------------------------------------------------------------------ schedule change log
def change_log(S, path):
    rr = route_rows(S)
    wb = Workbook()
    ws = wb.active
    ws.title = "Log"
    ws.append(["Ref", "Season", "Route", "Change", "From", "To", "Effective from", "Entered", "Note"])
    for c in ws[1]:
        c.font = Font(bold=True)
    ref = 0
    rows = []
    # IS24 changes against IS23 (history)
    for b, st_ in [("VRN", "OLB"), ("BLQ", "TPS"), ("LIN", "CIY")]:
        ref += 1
        rows.append(["SC-24-%03d" % ref, "IS24", "%s-%s" % (b, st_), "new route", "", "twice daily",
                     "31/03/2024", "06/10/2023", ""])
    for b, st_ in [("TRN", "BRI"), ("VCE", "TPS")]:
        ref += 1
        rows.append(["SC-24-%03d" % ref, "IS24", "%s-%s" % (b, st_), "frequency", "daily", "twice daily",
                     "31/03/2024", "20/10/2023", ""])
    ref = 0
    for r in rr:
        if r["kind"] != "past":
            continue
        ref += 1
        rows.append(["SC-25-%03d" % ref, "IS25", "%s-%s" % (r["base"], r["station"]), "gauge",
                     "%s (%d)" % (TYPE_LONG[r["t24"]], W.SEATS[r["t24"]]),
                     "%s (%d)" % (TYPE_LONG[r["t25"]], W.SEATS[r["t25"]]), "30/03/2025", "14/10/2024", ""])
    ref += 1
    rows.append(["SC-25-%03d" % ref, "IS25", "all", "flight numbers", "three-digit", "four-digit", "30/03/2025",
                 "04/11/2024", "numbers re-issued with the new reservation and revenue management systems; see "
                               "sheet IS25 renumbering"])
    ref = 0
    for r in rr:
        if r["kind"] != "target":
            continue
        ref += 1
        eff = "01/06/2026" if r["switch"] else "29/03/2026"
        ent = "21/11/2025" if r["switch"] else "24/10/2025"
        rows.append(["SC-26-%03d" % ref, "IS26", "LIN-%s" % r["station"], "gauge", "A319 (144)", "A321neo (232)", eff,
                     ent, "A319 to 31/05/2026" if r["switch"] else ""])
    for row in rows:
        ws.append(row)
    for col, w in zip("ABCDEFGHI", [11, 7, 9, 13, 13, 14, 13, 11, 70]):
        ws.column_dimensions[col].width = w
    ws2 = wb.create_sheet("IS25 renumbering")
    ws2.append(["Route", "Direction", "Rotation", "IS24 flight", "IS25 flight", "Departure (local)"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for b, st_, kind in route_list():
        n_new, n_old = BLOCKS[(b, st_)]
        freq = [r["freq"] for r in rr if r["base"] == b and r["station"] == st_][0]
        for outbound in (True, False):
            times = W.OUT_TIMES[freq] if outbound else W.RET_TIMES[freq]
            for k in range(freq):
                ws2.append(["%s-%s" % (b, st_), "%s-%s" % ((b, st_) if outbound else (st_, b)), k + 1,
                            n_old + 2 * k + (1 if outbound else 2), n_new + 2 * k + (1 if outbound else 2), times[k]])
    fix_workbook_props(wb, "Schedule change log", "Network Planning", dt.datetime(2025, 11, 21, 16, 12))
    wb.save(path)
    normalise_zip(path)


# ------------------------------------------------------------------ fleet register and route master
def fleet_register(path, regs):
    rng = sim.stream("fleet")
    rows = []
    # A319 exits: the redelivery slots agreed with the lessor, Monday to Thursday, through the winter
    exits, d = [], dt.date(2025, 11, 4)
    for k in range(10):
        exits.append(d)
        d = d + dt.timedelta(days=int(rng.integers(9, 21)))
        while d.weekday() > 3:
            d += dt.timedelta(days=1)
    for k, reg in enumerate(regs["319"]):
        exit_ = exits[k] if k < len(exits) else None
        status = "retired" if exit_ is not None and exit_ <= dt.date(2025, 11, 30) else "in service"
        y, m = 2002 + int(rng.integers(0, 5)), 1 + int(rng.integers(0, 12))
        rows.append([reg, "A319-112", 144, "owned" if k % 3 else "leased", "%d-%02d" % (y, m),
                     exit_.isoformat() if exit_ else "", status])
    # the A320s left for the A321s of the summer 2025 gauge changes
    a320_exit = [dt.date(2025, 2, 17), dt.date(2025, 3, 3), dt.date(2025, 3, 12), dt.date(2025, 3, 20), dt.date(2025, 3, 26)]
    for reg, ex in zip(regs["320"], a320_exit):
        y, m = 2008 + int(rng.integers(0, 4)), 1 + int(rng.integers(0, 12))
        rows.append([reg, "A320-214", 180, "leased", "%d-%02d" % (y, m), ex.isoformat(), "returned"])
    for k, reg in enumerate(regs["32N"]):
        if k < 8:
            y, m = 2021 + int(rng.integers(0, 3)), 1 + int(rng.integers(0, 12))
            ins = "%d-%02d" % (y, m)
        else:
            ins = "2025-%02d" % (1 + int(rng.integers(0, 3)))
        rows.append([reg, "A320-251N", 186, "leased", ins, "", "in service"])
    for reg in regs["321"]:
        ins = "%d-%02d" % (2024 + int(rng.integers(0, 2)), [11, 12, 1, 2, 3][int(rng.integers(0, 5))])
        rows.append([reg, "A321-231", 220, "leased", ins, "", "in service"])
    deliv = [("I-NQAA", "2026-01-15"), ("I-NQAB", "2026-01-28"), ("I-NQAC", "2026-02-05"), ("I-NQAD", "2026-02-11"),
             ("I-NQAE", "2026-02-19"), ("I-NQAF", "2026-02-26"), ("I-NQAG", "2026-03-05"), ("I-NQAH", "2026-03-17"),
             ("I-NQAL", "2026-05-13"), ("I-NQAM", "2026-05-27")]
    for reg, dd in deliv:
        rows.append([reg, "A321-271NX", 232, "leased", dd, "", "on order"])
    regs_all = [r[0] for r in rows]
    assert len(regs_all) == len(set(regs_all)), "a registration is listed twice"
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["registration", "model", "seats", "ownership", "in_service", "exit", "status_20251130"])
        w.writerows(rows)


def route_master(S, path):
    rr = route_rows(S)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["route", "base", "base_name", "station", "station_name", "region", "great_circle_km",
                    "first_flown"])
        first = {("VRN", "OLB"): "IS24", ("BLQ", "TPS"): "IS24", ("LIN", "CIY"): "IS24"}
        for r in rr:
            w.writerow(["%s-%s" % (r["base"], r["station"]), r["base"], STATION_NAME[r["base"]], r["station"],
                        STATION_NAME[r["station"]], REGION[r["station"]], r["dist"],
                        first.get((r["base"], r["station"]), "before IS23")])


# ------------------------------------------------------------------ summer close-out summary
def closeout(S, path):
    st = [Paragraph("Navigli Air - Commercial", SMALL),
          Paragraph("Summer close-out by route: IS24 and IS25", H1),
          Paragraph("Issued 14/11/2025 by Revenue Management reporting. Scheduled departures as flown, both "
                    "directions of each route. Load factor is seats sold over seats. Passengers flown are from "
                    "departure control; the IS24 figures include lap infants (the definition before the November "
                    "2024 migration), the IS25 figures exclude them.", SMALL), Spacer(1, 3 * mm)]
    for season in ("S24", "S25"):
        legs = S.legs(season)
        agg = defaultdict(lambda: [0, 0, 0, 0])
        for L in legs:
            key = (L.base, L.station)
            a = agg[key]
            a[0] += L.n
            a[1] += int(L.seats.sum())
            a[2] += int(L.sold.sum())
            a[3] += int(L.flown.sum() + (L.inf.sum() if season == "S24" else 0))
        data = [["Route", "Departures", "Seats", "Seats sold", "Passengers flown", "Load factor"]]
        tot = [0, 0, 0, 0]
        for b, s_, kind in route_list():
            a = agg[(b, s_)]
            data.append(["%s-%s" % (b, s_), "{:,}".format(a[0]), "{:,}".format(a[1]), "{:,}".format(a[2]),
                         "{:,}".format(a[3]), "%.1f%%" % (100.0 * a[2] / a[1])])
            tot = [x + y for x, y in zip(tot, a)]
        data.append(["Network", "{:,}".format(tot[0]), "{:,}".format(tot[1]), "{:,}".format(tot[2]),
                     "{:,}".format(tot[3]), "%.1f%%" % (100.0 * tot[2] / tot[1])])
        label = "IS24 (31/03/2024 - 26/10/2024)" if season == "S24" else "IS25 (30/03/2025 - 25/10/2025)"
        st.append(Paragraph(label, H2))
        t = _table(data, [24 * mm, 22 * mm, 26 * mm, 26 * mm, 30 * mm, 22 * mm])
        st.append(t)
        if season == "S24":
            st.append(PageBreak())
    _pdf(path, st, "Summer close-out by route IS24 IS25", "Revenue Management reporting",
         when=dt.datetime(2025, 11, 14, 10, 0))


# ------------------------------------------------------------------ the data dictionary
def dictionary(path, refusal_span):
    st = [Paragraph("Navigli Air - Revenue Management Systems", SMALL),
          Paragraph("Data dictionary: analytics extracts", H1),
          Paragraph("Revision 5, 2 December 2025. Owner: Concetta Gaito, Revenue Management Systems.", SMALL)]

    def block(title, intro, fields):
        st.append(Paragraph(title, H2))
        st.append(Paragraph(intro, BODY))
        st.append(_table([["Field", "Definition"]] + [[a, Paragraph(b, SMALL)] for a, b in fields],
                         [32 * mm, 146 * mm], align_right_from=9))

    block("rm_booking_build_s25.csv",
          "One row per scheduled departure (flight number and local departure date) of summer 2025, every route, as "
          "held by the revenue management system at the extract.",
          [("extract_ts", "time of the extract, UTC"),
           ("leg_key", "departure key: local date (yymmdd), scheduled time (hhmm), flight number"),
           ("flight_no", "four-digit flight number"),
           ("dep_date", "departure date, local"),
           ("std_local", "scheduled departure time, local"),
           ("orig, dest", "IATA airport codes"),
           ("equip", "aircraft type (319 A319, 320 A320, 32N A320neo, 321 A321, 32Q A321neo)"),
           ("cap", "seats on sale; domestic flights are not overbooked"),
           ("sold_out_utc", "time the last seat was sold, UTC; blank if the departure did not sell out"),
           ("bkd_dNNN", "seats sold (net bookings on hand) at the end of the local day NNN days before the departure "
                        "date; bkd_d000 is the count at departure. Snapshots: daily from day 14, every two days from "
                        "day 60, every five days from day 120, then days 140, 160, 180, 210, 240, 270, 300 and 330")])
    block("ym3_leg_history_s24.sqlite",
          "The YM3 yield management archive for summer 2024, frozen at the migration to the new revenue management "
          "system in November 2024. Tables legs, bkg and flown join on leg_id; ym3_extract holds the archive notes.",
          [("legs.flt", "summer 2024 flight number (three digits; the summer 2025 numbers are in the schedule change "
                        "log)"),
           ("legs.dep_dt, legs.std", "departure date (DD/MM/YYYY) and scheduled time (HHMM), local"),
           ("legs.typ, legs.cap", "aircraft type and seats on sale"),
           ("bkg.dcp, bkg.bkd", "seats sold at the YM3 data collection point dcp (days before departure, end of "
                                "day): 330, 270, 210, 180, 150, 120, 90, 75, 60, 45, 35, 28, 21, 14, 10, 7, 3, 1, 0"),
           ("bkg.avl", "availability after the snapshot: O seats left to sell, C none left"),
           ("flown.pax, flown.inf", "passengers flown on the leg, lap infants included; of which lap infants")])
    block("dcs_flight_close_s25.csv",
          "Departure control close-out, one row per operated departure of summer 2025.",
          [("pax_booked", "seats sold at departure"),
           ("pax_flown", "passengers flown, lap infants excluded"),
           ("inf_flown", "lap infants flown"),
           ("no_show", "passengers holding a seat who did not fly"),
           ("tail", "aircraft registration")])
    block("booking_engine_refusals_s25.csv",
          "Records written by the booking engines (website WEB, app APP, contact centre CC, agency feed GDS) when a "
          "request for a departure could not be sold, for summer 2025 departures; the requests in this extract were "
          "logged %s." % refusal_span,
          [("event_id", "record number"),
           ("logged_utc", "time the record was written, UTC"),
           ("channel", "WEB, APP, CC or GDS"),
           ("session_ref", "the engine's session reference (website and app), call reference (contact centre) or "
                           "booking reference (agency feed)"),
           ("flight_no, dep_date", "the departure requested (four-digit number, local date)"),
           ("orig, dest", "IATA airport codes of the journey requested"),
           ("fare_family", "fare family requested: LIGHT, CLASSIC or FLEX"),
           ("pax", "passengers in the request"),
           ("result", "NO_AVAIL")])
    _pdf(path, st, "Data dictionary analytics extracts rev 5", "Concetta Gaito", when=dt.datetime(2025, 12, 2, 11, 0))


# ------------------------------------------------------------------ distractors
def fare_report(S, path):
    """Revenue Management's monthly fare report: average net fare per passenger flown, by route and month."""
    rng = sim.stream("fares")
    wb = Workbook()
    ws = wb.active
    ws.title = "IS25 fares"
    ws.append(["Navigli Air - Revenue Management - fare report IS25 (EUR, net of taxes and airport charges)"])
    ws["A1"].font = Font(bold=True)
    ws.append(["Average net fare per passenger flown, by route (both directions) and month of departure"])
    ws.append([])
    months = ["Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct"]
    ws.append(["Route"] + months + ["Season"])
    for c in ws[4]:
        c.font = Font(bold=True)
    for b, st_, kind in route_list():
        dist = W.DISTANCE_KM[(b, st_)]
        base_fare = 38 + 0.072 * dist + (14 if kind == "target" else (9 if kind == "past" else 0))
        vals = []
        for m, mf in zip(range(3, 11), [0.86, 0.90, 0.95, 1.06, 1.22, 1.34, 1.04, 0.91]):
            vals.append(round(base_fare * mf * (1 + rng.normal(0, 0.025)), 2))
        ws.append(["%s-%s" % (b, st_)] + vals + [round(sum(vals[1:8]) / 7.0 * (1 + rng.normal(0, 0.004)), 2)])
    fix_workbook_props(wb, "Fare report IS25", "Revenue Management", dt.datetime(2025, 11, 10, 12, 30))
    wb.save(path)
    normalise_zip(path)


def competitor_seats(S, path):
    rr = route_rows(S)
    rng = sim.stream("competitors")
    other = {"CTA": 9800, "PMO": 8900, "NAP": 7400, "BRI": 6100, "BDS": 3900, "SUF": 4300, "REG": 1500,
             "CAG": 5600, "TPS": 900, "AHO": 1300, "CRV": 600, "PSR": 800, "CIY": 0, "OLB": 4100}
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["airport_pair", "season", "our_weekly_seats", "other_carriers_weekly_seats", "basis"])
        for r in rr:
            pair = "%s-%s" % (r["base"], r["station"])
            for season, seats in (("IS25", r["seats25"]), ("IS26", r["seats26"])):
                ours = int(round(seats / 30.0))
                base = other[r["station"]] * (1.0 if r["base"] == "LIN" else 0.35)
                oth = int(round(base * (1 + rng.normal(0, 0.03)) / 6.0) * 6) if base > 0 else 0
                w.writerow([pair, season, ours, oth, "filed schedules, typical week in July"])


def export_manifest(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["file", "owner", "system", "prepared", "covers", "notes"])
        w.writerows(rows)
