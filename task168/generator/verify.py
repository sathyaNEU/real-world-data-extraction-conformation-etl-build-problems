"""verify.py: the independent verifier. Reads only the shipped bundle, on its own code path (pandas, pdfminer),
and recomputes every graded figure and the rivals the write-up names; compares with generator/figures.json.

    python3 task168/generator/verify.py <target dir>
"""
import json
import math
import os
import sqlite3
import sys

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
NINE = {"CTA", "PMO", "NAP", "BRI", "BDS", "SUF", "REG", "CAG", "TPS"}
OUT = []


def ok(label, cond, detail=""):
    OUT.append(bool(cond))
    print("%s  %-62s %s" % ("PASS" if cond else "FAIL", label, detail))


def pdf_text(path):
    from pdfminer.high_level import extract_text
    return " ".join(extract_text(path).split())


def main(target):
    figs = json.load(open(os.path.join(HERE, "figures.json")))
    rm = pd.read_csv(os.path.join(target, "rm_booking_build_s25.csv"), dtype={"sold_out_utc": str})
    dcs = pd.read_csv(os.path.join(target, "dcs_flight_close_s25.csv"))
    ref = pd.read_csv(os.path.join(target, "booking_engine_refusals_s25.csv"), dtype={"session_ref": str})
    cols = [c for c in rm.columns if c.startswith("bkd_d")]
    t = np.array([int(c[5:]) for c in cols], float)
    ok("format csv: RM build reads back", len(rm) == 28140 and len(cols) == 58, "%d rows, %d snapshots" % (len(rm), len(cols)))

    rm["dep"] = pd.to_datetime(rm["dep_date"])
    rm["nine"] = ((rm["orig"] == "LIN") & rm["dest"].isin(NINE)) | ((rm["dest"] == "LIN") & rm["orig"].isin(NINE))
    rm["sold"] = rm["sold_out_utc"].notna() & (rm["sold_out_utc"] != "")
    B = rm[cols].to_numpy(float)
    fin = B[:, -1]
    # month build from the open departures off the nine routes, mean of shares
    refmask = (~rm["nine"]) & (~rm["sold"]) & (fin > 0)
    share = B / np.where(fin > 0, fin, 1)[:, None]
    month = rm["dep"].dt.month.to_numpy()
    curves = {m: share[refmask.to_numpy() & (month == m)].mean(axis=0) for m in np.unique(month)}
    # days from the sell-out to the end of the local departure date
    so = pd.to_datetime(rm["sold_out_utc"], utc=True, errors="coerce").dt.tz_convert("Europe/Rome")
    end = (rm["dep"] + pd.Timedelta(days=1)).dt.tz_localize("Europe/Rome")
    x = ((end - so).dt.total_seconds() / 86400.0).to_numpy()
    cap = rm["cap"].to_numpy(float)
    D = fin.copy()
    for k in np.nonzero(rm["sold"].to_numpy() & rm["nine"].to_numpy())[0]:
        c = curves[month[k]]
        D[k] = cap[k] / np.interp(x[k], t[::-1], c[::-1])
    # summer 2026 seats from the capacity plan workbook: the A321neo, the two routes on the A319 to 31 May
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(target, "summer2026_capacity_plan_v3.xlsx"))
    routes = pd.DataFrame(list(wb["Routes"].iter_rows(min_row=2, values_only=True)),
                          columns=[c.value for c in wb["Routes"][1]]).dropna(subset=["Base"])
    ok("format xlsx: capacity plan reads back", len(routes) == 47, "%d routes" % len(routes))
    seat26 = {}
    switch = {}
    for _, r in routes[routes["Base"] == "LIN"].iterrows():
        if r["Station"] in NINE:
            seat26[r["Station"]] = int(r["Seats/dep S26"])
            if str(r["S26 type from"]).startswith("01/06/2026"):
                switch[r["Station"]] = pd.Timestamp("2026-06-01")
    added = int([row for row in wb["Gauge changes"].iter_rows(values_only=True) if row[0] == "Nine Linate routes"][0][4])
    st = np.where(rm["orig"] == "LIN", rm["dest"], rm["orig"])
    d26 = rm["dep"] + pd.Timedelta(days=364)
    s26 = np.array([(144 if (s in switch and d < switch[s]) else seat26.get(s, 0)) for s, d in zip(st, d26)], float)
    # no-show share by route direction, summer 2025
    m = rm.merge(dcs, on=["flight_no", "dep_date"], suffixes=("", "_dcs"))
    ok("join RM build to DCS close-out one to one", len(m) == len(rm), "%d" % len(m))
    nm = m[m["nine"]]
    ns = nm.groupby(["orig", "dest"]).apply(lambda g: 1 - g["pax_flown"].sum() / g["pax_booked"].sum(),
                                           include_groups=False)
    m["D"] = D
    m["s26"] = s26
    m["ns"] = [ns.get((o, d), 0.0) for o, d in zip(m["orig"], m["dest"])]
    nm = m[m["nine"]]
    base = int(nm["pax_flown"].sum())
    fc = float((np.minimum(nm["D"], nm["s26"]) * (1 - nm["ns"])).sum())
    gain = fc - base
    call = int(math.floor(gain / 25000 + 0.5) * 25000)
    ok("call: additional passengers to the nearest 25,000", call == figs["call_additional_passengers"],
       "%.0f -> %d" % (gain, call))
    edge = 12500 - abs(gain - round(gain / 25000) * 25000)
    ok("call clears its bin edge by a quarter bin or more", edge >= 6250, "%.0f from the nearest edge" % edge)
    ok("component: summer 2025 passengers carried", base == figs["s25_passengers_carried"], str(base))
    ok("component: sold-out departures of the nine", int(nm["sold"].sum()) == figs["s25_departures_sold_out"]
       and len(nm) == figs["s25_departures_nine"], "%d of %d" % (nm["sold"].sum(), len(nm)))
    ok("seats added from the capacity plan", added == figs["s26_seats_added"], str(added))

    # rivals the write-up names
    pooled = share[refmask.to_numpy()].mean(axis=0)
    own = share[(rm["nine"] & ~rm["sold"] & (fin > 0)).to_numpy()].mean(axis=0)

    def with_curve(c):
        Dv = fin.copy()
        for k in np.nonzero(rm["sold"].to_numpy() & rm["nine"].to_numpy())[0]:
            Dv[k] = cap[k] / np.interp(x[k], t[::-1], c[::-1])
        sub = m["nine"].to_numpy()
        return float((np.minimum(Dv[sub], s26[sub]) * (1 - m["ns"].to_numpy()[sub])).sum()) - base
    pv, ov = with_curve(pooled), with_curve(own)
    ok("rival: pooled build about %s" % format(figs["rival_pooled_about"], ","), round(pv, -4) == figs["rival_pooled_about"],
       "%.0f" % pv)
    ok("rival: the nine routes' own open departures about %s" % format(figs["rival_own_open_about"], ","),
       round(ov, -4) == figs["rival_own_open_about"], "%.0f" % ov)
    # the 2025 gauge changes' passengers per added seat: YM3 summer 2024 against DCS summer 2025
    con = sqlite3.connect(os.path.join(target, "ym3_leg_history_s24.sqlite"))
    legs = pd.read_sql("select l.leg_id, l.orig, l.dest, l.typ, l.cap, f.pax, f.inf from legs l join flown f using (leg_id)", con)
    nb = pd.read_sql("select count(*) n from bkg", con)["n"][0]
    con.close()
    ok("format sqlite: YM3 archive reads back", len(legs) == 28140 and nb == 28140 * 19, "%d legs, %d snapshots" % (len(legs), nb))
    log = pd.DataFrame(list(load_workbook(os.path.join(target, "schedule_change_log.xlsx"))["Log"].iter_rows(min_row=2, values_only=True)),
                       columns=["ref", "season", "route", "change", "from", "to", "eff", "entered", "note"])
    past = set(log[(log["season"] == "IS25") & (log["change"] == "gauge")]["route"])
    bases = {"LIN", "TRN", "VRN", "VCE", "BLQ"}

    def route(o, d):
        return "%s-%s" % ((o, d) if o in bases else (d, o))
    legs["route"] = [route(o, d) for o, d in zip(legs["orig"], legs["dest"])]
    m["route"] = [route(o, d) for o, d in zip(m["orig"], m["dest"])]
    p24 = legs[legs["route"].isin(past)]
    p25 = m[m["route"].isin(past)]
    per_seat = ((p25["pax_flown"].sum() - (p24["pax"] - p24["inf"]).sum()) /
                (p25["cap"].sum() - p24["cap"].sum()))
    ok("rival: %.2f extra passengers per extra seat" % figs["rival_per_seat"], round(per_seat, 2) == figs["rival_per_seat"],
       "%.4f over %d routes" % (per_seat, len(past)))
    ok("rival: per seat onto the seats added about %s" % format(figs["rival_per_seat_total_about"], ","),
       round(per_seat * added, -4) == figs["rival_per_seat_total_about"], "%.0f" % (per_seat * added))

    # the holdout the write-up states: open departures off the nine routes censored at 85 per cent
    errs_m, errs_p = {}, {}
    hm = (refmask & (pd.Series(fin) >= 40)).to_numpy()
    rr = np.nonzero(hm)[0]
    tot = {}
    sm = {mm: share[hm & (month == mm)].sum(axis=0) for mm in curves}
    nmn = {mm: int((hm & (month == mm)).sum()) for mm in curves}
    sp, npool = share[hm].sum(axis=0), int(hm.sum())
    for k in rr:
        mm = month[k]
        cm = (sm[mm] - share[k]) / (nmn[mm] - 1)
        cpool = (sp - share[k]) / (npool - 1)
        c85 = 0.85 * fin[k]
        j = int(np.nonzero(B[k] >= c85)[0][0])
        if j == 0:
            continue
        b0, b1 = B[k, j - 1], B[k, j]
        xx = t[j - 1] - (t[j - 1] - t[j]) * ((c85 - b0) / (b1 - b0) if b1 > b0 else 1.0)
        a = tot.setdefault(mm, [0.0, 0.0, 0.0])
        a[0] += c85 / np.interp(xx, t[::-1], cm[::-1])
        a[1] += c85 / np.interp(xx, t[::-1], cpool[::-1])
        a[2] += fin[k]
    me = max(abs(a[0] / a[2] - 1) for a in tot.values())
    pj, pa = tot[7][1] / tot[7][2] - 1, tot[8][1] / tot[8][2] - 1
    ok("holdout: month builds within %d per cent in every month" % figs["holdout_month_within_pct"],
       me * 100 <= figs["holdout_month_within_pct"], "max %.2f per cent" % (100 * me))
    ok("holdout: pooled build over %d per cent high in July and August" % figs["holdout_pooled_peak_over_pct"],
       min(pj, pa) * 100 > figs["holdout_pooled_peak_over_pct"], "Jul %.2f, Aug %.2f per cent" % (100 * pj, 100 * pa))

    # the ask: passengers turned away because the flight was full
    r9 = ref[((ref["orig"] == "LIN") & ref["dest"].isin(NINE)) | ((ref["dest"] == "LIN") & ref["orig"].isin(NINE))]
    r9 = r9.merge(rm[["flight_no", "dep_date", "sold_out_utc"]], on=["flight_no", "dep_date"], how="left")
    full = r9[r9["sold_out_utc"].notna() & (pd.to_datetime(r9["logged_utc"], utc=True) >= pd.to_datetime(r9["sold_out_utc"], utc=True))]
    web = full[full["channel"].isin(["WEB", "APP"])].drop_duplicates(["session_ref", "flight_no", "dep_date"])
    rest = full[full["channel"].isin(["CC", "GDS"])]
    cnt = pd.concat([web, rest]).groupby(["orig", "dest"])["pax"].sum()
    got = {"%s-%s" % k: int(v) for k, v in cnt.items()}
    ok("ask: 18 route directions match the figures", got == figs["turned_away"], "total %d" % sum(got.values()))
    ok("format csv: refusals log reads back", len(ref) > 25000, "%d rows" % len(ref))

    # the documents carrying the filed sentences, through pdfminer, with positive controls on the same path
    basis = pdf_text(os.path.join(target, "planning_basis_NP-PB-26S_issue2.pdf"))
    for frag in ("Lap infants are recorded separately and are not passengers carried",
                 "we do not overbook domestic flights", "wanted to fly the nine routes",
                 "repeats summer 2025's departures", "additional passengers carried on the nine routes"):
        ok("format pdf: planning basis carries '%s'" % frag[:44], frag in basis, "")
    ok("positive control: a sentence the basis does not carry is not found", "carried plus" not in basis, "")
    dic = pdf_text(os.path.join(target, "data_dictionary_rev5.pdf"))
    ok("format pdf: dictionary defines sold_out_utc and the snapshots", "time the last seat was sold" in dic
       and "end of the local day" in dic, "")
    clo = pdf_text(os.path.join(target, "summer_closeout_IS24_IS25.pdf"))
    net25 = int(dcs["pax_flown"].sum())
    ok("format pdf: close-out IS25 network passengers equal the DCS total", format(net25, ",") in clo, format(net25, ","))
    print("\n%d checks, %d failed" % (len(OUT), OUT.count(False)))
    return 0 if all(OUT) else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(TASK, "target")))
