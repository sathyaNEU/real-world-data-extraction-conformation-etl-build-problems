"""The goldens, generator/figures.json and submission.md, every figure computed from target/."""
import datetime as dt
import json
import math
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import calc as C
import rivals as R
from fileio import fix_workbook_props, normalise_zip

ROUTE_ORDER = ["CTA", "PMO", "NAP", "BRI", "BDS", "SUF", "REG", "CAG", "TPS"]
NAME = {"CTA": "Catania", "PMO": "Palermo", "NAP": "Napoli", "BRI": "Bari", "BDS": "Brindisi",
        "SUF": "Lamezia Terme", "REG": "Reggio Calabria", "CAG": "Cagliari", "TPS": "Trapani"}
CALL_UNIT = 25000


def r_to(v, u):
    return int(math.floor(v / u + 0.5) * u)


def compute(target):
    g = C.golden(target)
    fam = R.family(target, g)
    rg = R.rungs(target, g)
    bt = R.backtest(target, g)
    pc = R.pseudo_censor(target, g)
    ta = C.turnaways(target, g["rows"], "golden")
    rows, dcs = g["rows"], g["dcs"]
    nine = [r for r in rows if C.is_nine(r["orig"], r["dest"])]
    sold = sum(1 for r in nine if r["so"])
    resp = R.per_seat_response(target, rows, dcs)
    over = sum(1 for r in nine if C.demand(r, g["curves"][r["dep"].month], g["dcps"]) > 232)
    return dict(g=g, fam=fam, rg=rg, bt=bt, pc=pc, ta=ta, n_nine=len(nine), n_sold=sold, resp=resp, n_over_232=over)


def figures(c):
    g, rg, pc, ta = c["g"], c["rg"], c["pc"], c["ta"]
    month_err = max(abs(v["month"]) for v in pc.values())
    f = {
        "call_additional_passengers": r_to(g["gain"], CALL_UNIT),
        "call_rounding": CALL_UNIT,
        "s25_passengers_carried": g["base"],
        "s25_departures_nine": c["n_nine"],
        "s25_departures_sold_out": c["n_sold"],
        "s26_seats_added": rg["added_seats"],
        "a321neo_seats": 232,
        "a319_seats": 144,
        "replay_days": 364,
        "rival_per_seat": round(rg["per_seat"], 2),
        "past_seat_increase_min_pct": int(round(100 * min(c["resp"]["seats25"][k] / c["resp"]["seats24"][k] - 1
                                                         for k in c["resp"]["routes"]))),
        "past_seat_increase_max_pct": int(round(100 * max(c["resp"]["seats25"][k] / c["resp"]["seats24"][k] - 1
                                                         for k in c["resp"]["routes"]))),
        "nine_departures_demand_over_232": c["n_over_232"],
        "rival_per_seat_total_about": r_to(rg["rung0_per_seat"], 10000),
        "rival_pooled_about": r_to(rg["rung2_pooled_network"], 10000),
        "rival_own_open_about": r_to(rg["rung2b_pooled_nine_own_open"], 10000),
        "holdout_censor_share_pct": 85,
        "holdout_month_within_pct": int(math.ceil(month_err * 100)),
        "holdout_pooled_peak_over_pct": 10 if min(pc[7]["pooled"], pc[8]["pooled"]) > 0.10 else None,
        "turned_away": {"%s-%s" % k: v for k, v in sorted(ta.items())},
    }
    for st in ROUTE_ORDER:
        for o, d in (("LIN", st), (st, "LIN")):
            f["turned_away_%s_%s" % (o, d)] = ta[(o, d)]
    return f


# ------------------------------------------------------------------ the forecast workbook
THIN = Side(style="thin", color="9AA5B1")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
BODY = Font(size=10)


def _head(ws, row, values, widths=None):
    for k, v in enumerate(values, start=1):
        cell = ws.cell(row=row, column=k, value=v)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if k > 1 else "left")
    if widths:
        for k, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(k)].width = w


def per_route(c):
    g = c["g"]
    rows, dcs, dcps = g["rows"], g["dcs"], g["dcps"]
    ns = C.noshow_by_leg(rows, dcs)
    cm = g["curves"]
    out = {}
    for r in rows:
        if not C.is_nine(r["orig"], r["dest"]):
            continue
        st = C.station(r["orig"], r["dest"])
        a = out.setdefault(st, dict(dep=0, sold=0, base=0, fc=0.0))
        a["dep"] += 1
        a["sold"] += 1 if r["so"] else 0
        a["base"] += int(dcs[(r["flt"], r["dep"].isoformat())]["pax_flown"])
        D = C.demand(r, cm[r["dep"].month], dcps)
        a["fc"] += min(D, C.seats_s26(r)) * (1 - ns[(r["orig"], r["dest"])])
    return out


def workbook(c, f, path):
    g = c["g"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Nine Linate routes: about %s more passengers carried in summer 2026 than in summer 2025" % (
        "{:,}".format(f["call_additional_passengers"]))
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Network Planning, summer 2026 gauge change (A319 to A321neo). Figures as at 12/12/2025."
    ws["A2"].font = Font(italic=True, size=9, color="555555")
    _head(ws, 4, ["", "Value", "Basis"], [72, 14, 70])
    lines = [
        ("Additional passengers carried, summer 2026 over summer 2025 (nearest 25,000)",
         f["call_additional_passengers"], "forecast; see By route and Booking build"),
        ("Summer 2025 passengers carried", f["s25_passengers_carried"],
         "DCS close-out, lap infants and no-shows excluded (NP-PB-26S section 4)"),
        ("Summer 2025 departures, both directions", f["s25_departures_nine"], "RM booking build extract 01/12/2025"),
        ("  of which sold out", f["s25_departures_sold_out"], "sell-out time recorded by RM"),
        ("Summer 2026 seats added", f["s26_seats_added"], "capacity plan v3, 21/11/2025"),
    ]
    for k, (a, b, cc) in enumerate(lines, start=5):
        ws.cell(row=k, column=1, value=a).font = Font(bold=(k == 5), size=10)
        v = ws.cell(row=k, column=2, value=b)
        v.number_format = "#,##0"
        v.font = Font(bold=(k == 5), size=10)
        ws.cell(row=k, column=3, value=cc).font = Font(size=9, color="444444")
    ws["A11"] = ("Each sold-out departure's demand is its seats over the share of final bookings that departures of "
                 "the same month which never sold out (the airline's other domestic routes) held at the time it sold "
                 "out. Summer 2026 re-flies every summer 2025 departure 364 days later at its planned seats, bookings "
                 "capped at the seats, less the route direction's summer 2025 no-show share.")
    ws["A11"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A11:C13")
    ws.row_dimensions[11].height = 30
    ws.freeze_panes = "A5"

    ws = wb.create_sheet("By route")
    _head(ws, 1, ["Route", "Departures S25", "Sold out", "Passengers carried S25", "Passengers carried S26 (est.)",
                  "Change (est.)"], [22, 14, 11, 20, 22, 15])
    pr = per_route(c)
    k = 2
    for st in ROUTE_ORDER:
        a = pr[st]
        vals = ["Linate - %s" % NAME[st], a["dep"], a["sold"], a["base"], r_to(a["fc"], 1000),
                r_to(a["fc"] - a["base"], 1000)]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row=k, column=j, value=v)
            cell.font = BODY
            if j > 1:
                cell.number_format = "#,##0"
        k += 1
    tot = ["Nine routes", sum(pr[s]["dep"] for s in pr), sum(pr[s]["sold"] for s in pr), g["base"],
           r_to(g["forecast"], 1000), r_to(g["gain"], 1000)]
    for j, v in enumerate(tot, start=1):
        cell = ws.cell(row=k, column=j, value=v)
        cell.font = Font(bold=True, size=10)
        cell.border = Border(top=THIN)
        if j > 1:
            cell.number_format = "#,##0"
    ws.cell(row=k + 2, column=1, value="Estimates to the nearest 1,000 for working; the network figure is stated to the "
            "nearest 25,000 on the Summary. Rows may not add to the total because of rounding.").font = Font(
        italic=True, size=9, color="555555")
    ws.freeze_panes = "B2"

    ws = wb.create_sheet("Booking build")
    cm, ncur, dcps = g["curves"], g["n_curve"], g["dcps"]
    show = [120, 90, 60, 45, 30, 21, 14, 7, 3]
    _head(ws, 1, ["Departure month", "Departures in the build"] + ["%d days out" % t for t in show],
          [16, 13] + [10] * len(show))
    k = 2
    for m in sorted(cm):
        ws.cell(row=k, column=1, value=dt.date(2025, m, 1).strftime("%B")).font = BODY
        ws.cell(row=k, column=2, value=ncur[m]).number_format = "#,##0"
        for j, t in enumerate(show, start=3):
            cell = ws.cell(row=k, column=j, value=round(float(C.interp(cm[m], dcps, t)), 3))
            cell.number_format = "0.0%"
            cell.font = BODY
        k += 1
    ws.cell(row=k + 1, column=1, value="Share of final bookings on hand at the end of the day, mean over summer 2025 "
            "departures off the nine routes that never sold out (rm_booking_build_s25.csv).").font = Font(
        italic=True, size=9, color="555555")
    ws.freeze_panes = "C2"

    ws = wb.create_sheet("Turned away")
    _head(ws, 1, ["Route", "Direction", "Passengers turned away"], [22, 14, 20])
    k = 2
    for st in ROUTE_ORDER:
        for o, d in (("LIN", st), (st, "LIN")):
            ws.cell(row=k, column=1, value="Linate - %s" % NAME[st]).font = BODY
            ws.cell(row=k, column=2, value="%s-%s" % (o, d)).font = BODY
            cell = ws.cell(row=k, column=3, value=f["turned_away_%s_%s" % (o, d)])
            cell.number_format = "#,##0"
            cell.font = BODY
            k += 1
    ws.cell(row=k + 1, column=1, value="Summer 2025 departures. A passenger is turned away because the flight was full when "
            "the request was logged at or after the departure sold out. A website or app session asking again for "
            "the same departure is one party; the agency feed writes one record per passenger name; the contact "
            "centre one record per call.").alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=k + 1, start_column=1, end_row=k + 3, end_column=3)
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Notes")
    notes = [
        ("Scope", "Milano Linate to and from Catania, Palermo, Napoli, Bari, Brindisi, Lamezia Terme, Reggio Calabria, "
                  "Cagliari and Trapani; both directions (planning basis NP-PB-26S issue 2, section 2)."),
        ("Seasons", "Summer 2025 30/03/2025 - 25/10/2025 as flown; summer 2026 29/03/2026 - 24/10/2026 as scheduled."),
        ("Passengers carried", "Passengers flown at DCS close-out; lap infants and no-shows are not passengers carried."),
        ("Capacity", "Capacity plan v3: A321neo 232 seats; Brindisi and Lamezia Terme on the A319 (144) to 31/05/2026."),
        ("Market", "No growth (NP-PB-26S section 6)."),
        ("Rounding", "The additional passengers to the nearest 25,000; turned-away passengers whole."),
    ]
    _head(ws, 1, ["Item", "Note"], [20, 110])
    for k, (a, b) in enumerate(notes, start=2):
        ws.cell(row=k, column=1, value=a).font = Font(bold=True, size=10)
        ws.cell(row=k, column=2, value=b).alignment = Alignment(wrap_text=True)
    for sh in wb.worksheets:
        sh.page_setup.orientation = "landscape"
        sh.page_setup.paperSize = sh.PAPERSIZE_A4
        sh.sheet_properties.pageSetUpPr.fitToPage = True
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
    wb["Summary"].row_dimensions[12].height = 22
    wb["Summary"].row_dimensions[13].height = 22
    fix_workbook_props(wb, "Summer 2026 gauge change forecast", "Network Planning", dt.datetime(2025, 12, 12, 11, 40))
    wb.save(path)
    normalise_zip(path)


# ------------------------------------------------------------------ the one-slide chart
def deck(f, path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
    from pptx.util import Emu, Inches, Pt
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    top = max(((k, v) for k, v in f.items() if k.startswith("turned_away_")), key=lambda kv: kv[1])
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.9))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.text = "Summer 2025: Catania and Palermo turned away the most passengers once their flights were full"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    p2 = tb.text_frame.add_paragraph()
    p2.text = "Passengers turned away from summer 2025 departures because the flight was full, by route and direction"
    p2.font.size = Pt(13)
    p2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    cd = CategoryChartData()
    cd.categories = [NAME[st] for st in ROUTE_ORDER]
    cd.add_series("From Linate", [f["turned_away_LIN_%s" % st] for st in ROUTE_ORDER])
    cd.add_series("To Linate", [f["turned_away_%s_LIN" % st] for st in ROUTE_ORDER])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.35), Inches(12.3), Inches(5.4), cd)
    ch = gf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(12)
    colours = [RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0x8F, 0xA9, 0xC8)]
    for ser, col in zip(ch.series, colours):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = col
        ser.data_labels.show_value = True
        ser.data_labels.number_format = "#,##0"
        ser.data_labels.number_format_is_linked = False
        ser.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        ser.data_labels.font.size = Pt(10)
    ch.plots[0].gap_width = 60
    ch.plots[0].overlap = -10
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
    va.tick_labels.number_format = "#,##0"
    va.tick_labels.number_format_is_linked = False
    va.tick_labels.font.size = Pt(11)
    va.has_title = True
    va.axis_title.text_frame.text = "Passengers turned away"
    va.axis_title.text_frame.paragraphs[0].font.size = Pt(11)
    ch.category_axis.tick_labels.font.size = Pt(11)
    src = s.shapes.add_textbox(Inches(0.5), Inches(6.85), Inches(11.6), Inches(0.4))
    src.text_frame.word_wrap = True
    q = src.text_frame.paragraphs[0]
    q.text = ("Source: booking engines' refusals log (website, app, contact centre, agency feed), requests logged at or "
              "after the departure sold out (RM sell-out time). Network Planning, 12/12/2025.")
    q.font.size = Pt(10)
    q.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    num = s.shapes.add_textbox(Inches(12.4), Inches(6.85), Inches(0.5), Inches(0.4))
    num.text_frame.paragraphs[0].text = "1"
    num.text_frame.paragraphs[0].font.size = Pt(10)
    cp = prs.core_properties
    cp.title = "Summer 2025 passengers turned away"
    cp.author = "Network Planning"
    cp.last_modified_by = "Network Planning"
    cp.created = dt.datetime(2025, 12, 12, 11, 52)
    cp.modified = dt.datetime(2025, 12, 12, 11, 52)
    cp.revision = 1
    cp.subject = ""
    cp.keywords = ""
    cp.comments = ""
    prs.save(path)
    normalise_zip(path)


# ------------------------------------------------------------------ the write-up
def submission(f, path):
    ta = f
    fmt = "{:,}".format
    L = []
    L.append("## Tags\n")
    L.append("**Domain:** Urban & Regional Science (forecasting: an airline's passengers carried on nine Linate "
             "routes after a gauge change, each sold-out departure's demand completed from the booking build of its "
             "own season).")
    L.append("**Analytical objective:** Forecasting & Predictive Modeling.\n")
    L.append("## 1. Final Recommendation\n")
    L.append("**The nine Linate routes will carry about %s more passengers in summer 2026 than they carried in summer "
             "2025.** %s of their %s summer 2025 departures sold out, so each sold-out departure's demand is its seats "
             "over the share of final bookings that departures of the same month which never sold out held when it "
             "sold out; re-flown at summer 2026's seats less no-shows, that demand puts about %s more passengers into "
             "the %s seats added. Not the 2025 gauge changes' %s extra passengers per extra seat carried onto this "
             "one, about %s: those changes added %d to %d per cent to the seats of routes whose demand still filled "
             "most of them, while summer 2025's demand on the nine routes exceeds the A321neo's %d seats on almost no "
             "departure. Not a booking build pooled over the summer, about %s, nor one from the nine routes' own "
             "departures that never sold out, about %s: both are weighted to spring and autumn departures, whose "
             "passengers book a few weeks out, so they read every July and August departure as further from its "
             "final bookings than it was." % (
                 fmt(f["call_additional_passengers"]), fmt(f["s25_departures_sold_out"]),
                 fmt(f["s25_departures_nine"]), fmt(f["call_additional_passengers"]), fmt(f["s26_seats_added"]),
                 "%.2f" % f["rival_per_seat"], fmt(f["rival_per_seat_total_about"]),
                 f["past_seat_increase_min_pct"], f["past_seat_increase_max_pct"], f["a321neo_seats"],
                 fmt(f["rival_pooled_about"]), fmt(f["rival_own_open_about"])))
    L.append("")
    L.append("## 2. Critical Components\n")
    L.append("1. Summer 2025 passengers carried on the nine routes: **%s**." % fmt(f["s25_passengers_carried"]))
    L.append("2. Summer 2025 departures on the nine routes that sold out: **%s** of **%s**." % (
        fmt(f["s25_departures_sold_out"]), fmt(f["s25_departures_nine"])))
    L.append("")
    L.append("## 3. Step-by-Step Solution\n")
    steps = [
        "Took summer 2025 passengers carried from dcs_flight_close_s25.csv (pax_flown, lap infants and no-shows "
        "excluded per section 4 of planning_basis_NP-PB-26S_issue2.pdf): %s on the nine routes." % fmt(
            f["s25_passengers_carried"]),
        "In rm_booking_build_s25.csv, %s of the nine routes' %s summer 2025 departures carry a sold_out_utc, so their "
        "bookings stop at the seats." % (fmt(f["s25_departures_sold_out"]), fmt(f["s25_departures_nine"])),
        "From the summer 2025 departures off the nine routes that never sold out, built each departure month's mean "
        "share of final bookings on hand at every snapshot; censored at %d per cent of their final bookings, those "
        "departures are recovered within %d per cent in every month by their own month's build, while the build "
        "pooled over the summer puts their July and August totals more than %d per cent too high." % (
            f["holdout_censor_share_pct"], f["holdout_month_within_pct"], f["holdout_pooled_peak_over_pct"]),
        "Completed each sold-out departure on the nine routes as its seats over its month's share on hand at its "
        "sold_out_utc.",
        "Re-flew every summer 2025 departure %d days later at its seats in summer2026_capacity_plan_v3.xlsx (%d on "
        "the A321neo, %d on the A319 that Brindisi and Lamezia Terme keep to 31 May 2026), its bookings capped at the "
        "seats (section 5), less its route direction's summer 2025 no-show share from dcs_flight_close_s25.csv." % (
            f["replay_days"], f["a321neo_seats"], f["a319_seats"]),
        "Against summer 2025's %s passengers carried, the nine routes carry about %s more, to the nearest %s." % (
            fmt(f["s25_passengers_carried"]), fmt(f["call_additional_passengers"]), fmt(f["call_rounding"])),
        "Counted the turned-away passengers in booking_engine_refusals_s25.csv from requests logged at or after the "
        "departure's sold_out_utc, each website or app session once per departure at its pax, and every agency-feed "
        "and contact-centre record at its pax.",
    ]
    for k, s in enumerate(steps, start=1):
        L.append("%d. %s" % (k, s))
    L.append("")
    L.append("## 4. Deliverable Answers\n")
    L.append("### summer2026_upgauge_forecast.xlsx\n")
    L.append("1. Additional passengers carried on the nine routes, summer 2026 over summer 2025: %s." % fmt(
        f["call_additional_passengers"]))
    L.append("2. Passengers the booking engines turned away from summer 2025 departures because the flight was full, "
             "by route and direction:")
    for st in ROUTE_ORDER:
        for o, d in (("LIN", st), (st, "LIN")):
            L.append("   - %s-%s: %s" % (o, d, fmt(f["turned_away_%s_%s" % (o, d)])))
    L.append("")
    L.append("### summer2025_turnaways.pptx\n")
    L.append("1. One slide charting the counts in answer 2 above, a bar for each route and direction, each bar "
             "labelled with its count.")
    L.append("")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L))


def main(target, golden_dir, figures_path, submission_path):
    c = compute(target)
    f = figures(c)
    with open(figures_path, "w", encoding="utf-8") as fh:
        json.dump(f, fh, indent=1, sort_keys=True)
    os.makedirs(golden_dir, exist_ok=True)
    workbook(c, f, os.path.join(golden_dir, "summer2026_upgauge_forecast.xlsx"))
    deck(f, os.path.join(golden_dir, "summer2025_turnaways.pptx"))
    submission(f, submission_path)
    return c, f


if __name__ == "__main__":
    import sys
    here = os.path.dirname(os.path.abspath(__file__))
    task = os.path.dirname(here)
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(task, "target"), os.path.join(task, "golden"),
         os.path.join(here, "figures.json"), os.path.join(task, "submission.md"))
