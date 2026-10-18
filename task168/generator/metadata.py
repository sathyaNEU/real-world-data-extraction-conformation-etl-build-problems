"""metadata.json in the task folder: the portal's keys, the distractors and the provenance of every file."""
import csv
import json
import os
import sqlite3

DESC = {
    "rm_booking_build_s25.csv": "revenue management system, summer 2025 departure booking build (extract 1 December 2025)",
    "ym3_leg_history_s24.sqlite": "YM3 yield management archive, summer 2024 legs, snapshots and flown passengers",
    "dcs_flight_close_s25.csv": "departure control, summer 2025 flight close-out",
    "booking_engine_refusals_s25.csv": "booking engines, requests for summer 2025 departures that could not be sold",
    "summer2026_capacity_plan_v3.xlsx": "Network Planning, summer 2026 segment capacity plan version 3 (21 November 2025)",
    "planning_basis_NP-PB-26S_issue2.pdf": "Network Planning, summer 2026 planning basis NP-PB-26S issue 2 (28 November 2025)",
    "schedule_change_log.xlsx": "Network Planning, schedule change log IS24 to IS26 with the IS25 renumbering",
    "fleet_register.csv": "Fleet Management, fleet register as at 30 November 2025",
    "route_master.csv": "Network Planning, route master",
    "summer_closeout_IS24_IS25.pdf": "Revenue Management reporting, summer close-out by route IS24 and IS25 (14 November 2025)",
    "data_dictionary_rev5.pdf": "Revenue Management Systems, data dictionary for the analytics extracts, revision 5",
    "fare_report_IS25.xlsx": "Revenue Management, IS25 fare report",
    "competitor_seats_IS25_IS26.csv": "Network Planning, filed weekly seats by airport pair, IS25 and IS26",
    "export_manifest.csv": "Network Planning, export manifest for this pack (12 December 2025)",
}
DISTRACTORS = ["fare_report_IS25.xlsx", "competitor_seats_IS25_IS26.csv"]


def rows_of(path):
    if path.endswith(".csv"):
        with open(path, encoding="utf-8") as f:
            return sum(1 for _ in csv.reader(f)) - 1
    if path.endswith(".sqlite"):
        con = sqlite3.connect(path)
        n = sum(con.execute("select count(*) from %s" % t).fetchone()[0]
                for (t,) in con.execute("select name from sqlite_master where type='table'"))
        con.close()
        return n
    if path.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True)
        return sum(ws.max_row for ws in wb.worksheets)
    return None


def write(task_dir, n=None):
    target = os.path.join(task_dir, "target")
    files = []
    for f in sorted(os.listdir(target)):
        p = os.path.join(target, f)
        files.append({"path": f, "format": os.path.splitext(f)[1][1:], "bytes": os.path.getsize(p), "rows": rows_of(p),
                      "source": "Constructed for this task (seeded, deterministic): " + DESC[f],
                      "date": "2025-12-12", "license": "CC0-1.0 (original work, constructed for this task)"})
    largest = max(files, key=lambda x: x["rows"] or 0)
    meta = {
        "task": "task168",
        "title": ("How many more passengers an Italian domestic airline's nine Linate routes to the South and the "
                  "islands will carry in summer 2026 than in summer 2025, when their A319s give way to A321neos on "
                  "routes whose summer departures sell out"),
        "domain": "Urban & Regional Science",
        "client_domain": "Urban & Regional Science",
        "house_domain": "supply-chain-logistics",
        "subdomain": ("forecasting (transport demand: an airline's passengers carried after a gauge change, sold-out "
                      "departures' demand completed from the booking build of departures of the same season that "
                      "never sold out)"),
        "objective": "Forecasting & Predictive Modeling",
        "prompt_shape": ("a forecast figure (additional passengers carried, summer 2026 over summer 2025, nine routes), "
                         "with the passengers turned away from summer 2025 departures because the flight was full for "
                         "each of the 18 route directions, charted on one slide"),
        "as_of": "2025-12-12",
        "deliverables": ["summer2026_upgauge_forecast.xlsx", "summer2025_turnaways.pptx"],
        "distractor_files": DISTRACTORS,
        "field_terms": [],
        "stated_inputs": [],
        "input_gates": {"files": len(files), "formats": sorted({x["format"] for x in files}),
                        "largest_file": largest["path"], "largest_file_rows": largest["rows"]},
        "sources": [{
            "files": "all files under target/",
            "origin": ("constructed for this task by task168/generator/build.py (seeded, deterministic) around a "
                       "fictional Italian domestic airline (Navigli Air), its northern bases and its routes to real "
                       "Italian airports, in the layouts of an airline's revenue management extract, a legacy yield "
                       "management archive, a departure control close-out, booking engine logs and network planning "
                       "documents; no real organisation's records or person is depicted."),
            "date": "2025-12-12",
            "license": "CC0-1.0 (original work)"}],
        "created": "2026-10-18",
        "generator": "task168/generator/build.py",
        "files": files,
    }
    with open(os.path.join(task_dir, "metadata.json"), "w", encoding="utf-8") as fh:
        json.dump(meta, fh, indent=2, ensure_ascii=False)
        fh.write("\n")
    return meta
