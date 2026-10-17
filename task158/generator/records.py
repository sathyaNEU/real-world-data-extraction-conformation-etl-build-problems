"""task158 records: the files each of Slenk's systems exports, written from the world. Nothing here writes a truth
column: the rollup carries prompt tokens as the gateway counts them, never the tokens a model computed."""
from __future__ import annotations

import datetime as dt

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import common as K
import world as WD

F_ROLLUP = "fleet_rollup_hourly_2026-01-01_2026-12-06.csv"
F_BILLING = "docs_usage_by_org_daily_2026-09-01_2026-12-06.csv"
F_INTAKE = "q1_2027_capacity_intake.xlsx"
F_PLAN = "fleet_plan_2027Q1.xlsx"
F_LEDGER = "report_service_commitments_2026-12-04.xlsx"
F_NODES = "fleet_nodes_2026-12-06.csv"
F_JOBS = "fleetctl_jobs_2026-10-26_2026-12-06.jsonl"
F_CONSOLE = "ops_console_audit_2026-10-26_2026-12-06.csv"
F_BURN = "nx7_burnin_results.csv"
F_FABRIC = "nx7_fabric_results.csv"
F_HANDOVER = "nx7_handover_inventory_2026-12-04.csv"
F_INCIDENTS = "serving_incidents_2026H2.csv"
F_SPEC = "gpu_spec_register_2026-12.csv"

ROLLUP_COLS = ["hour_start_utc", "pool", "model", "gpus_in_service", "gpu_busy_s", "requests", "sessions_new",
               "prompt_tokens", "completion_tokens", "ttft_p50_ms", "ttft_p95_ms"]
BILLING_COLS = ["usage_date", "org_id", "org_name", "product", "channel", "requests", "sessions", "input_tokens",
                "output_tokens"]
NODE_COLS = ["node", "hall", "rack", "gpu_model", "gpus", "pool", "member_from", "member_to"]
CONSOLE_COLS = ["time (Europe/Amsterdam)", "user", "action", "object", "detail", "ticket"]
BURN_COLS = ["rack", "node", "slot", "gpu_serial", "run", "test", "started_utc", "finished_utc", "result", "note"]
FABRIC_COLS = ["rack", "node", "run", "test", "started_utc", "finished_utc", "algbw", "busbw"]
HANDOVER_COLS = ["rack", "position", "node", "slot", "gpu_serial"]
INCIDENT_COLS = ["incident", "opened_utc", "resolved_utc", "severity", "pools", "summary"]
SPEC_COLS = ["deployment", "model", "gpu", "gpus_per_node", "tensor_parallel", "replicas_at_peak", "max_num_seqs",
             "max_model_len", "kv_cache_dtype", "engine"]

ROLLUP_GAPS = {("chat-prod", dt.datetime(2026, 8, 14, 3, 0)), ("chat-prod", dt.datetime(2026, 8, 14, 4, 0)),
               ("chat-prod", dt.datetime(2026, 8, 14, 5, 0)), ("voice-prod", dt.datetime(2026, 3, 3, 11, 0))}


class Shaped:
    pass


def shape(W, NODES, M):
    """Integer counters the rollup and billing write, and the rollup's in-service samples."""
    S = Shaped()
    S.T = WD.pool_totals(W)
    hist = WD.random_cordons(NODES, np.random.default_rng(K.SEED + 35))
    S.in_service = WD.in_service(NODES, hist + M["stays"], W.hours)
    S.ttft = WD.ttft(W, S.T)
    return S


def model_tag(pool, t):
    """The model builds the pool served in the hour; both while the rolling upgrade of 11 and 12 May ran."""
    if pool in ("docs-prod", "batch-prod"):
        if t < K.LC2_START:
            return K.LC1_TAG
        if t + dt.timedelta(hours=1) <= K.LC2_END:
            return K.LC1_TAG + "+" + K.LC2_TAG
        return K.LC2_TAG
    return K.MODEL_TAG[pool]


def rollup_rows(W, S):
    for i, t in enumerate(W.hours):
        for p in K.POOLS:
            if (p, t) in ROLLUP_GAPS:
                continue
            d = S.T[p]
            yield [K.iso(t), p, model_tag(p, t), int(S.in_service[p][i]), int(round(W.busy[p][i])),
                   int(round(d["R"][i])), int(round(d["S"][i])), int(round(d["P"][i])), int(round(d["Q"][i])),
                   int(S.ttft[p][0][i]), int(S.ttft[p][1][i])]


def billing_rows(W):
    """Docs usage by organisation, channel and Amsterdam business day, 1 September to 6 December 2026."""
    O = W.orgs
    loc_dates = np.array([l.date() for l in W.local])
    first, last = dt.date(2026, 9, 1), dt.date(2026, 12, 6)
    rng = np.random.default_rng(K.SEED + 70)
    days = [first + dt.timedelta(days=i) for i in range((last - first).days + 1)]
    idx = {d: np.where(loc_dates == d)[0] for d in days}
    D = W.docs_org
    out = []
    for d in days:
        ii = idx[d]
        for o in range(O.n):
            r = D["R"][o, ii].sum()
            if r < 0.5:
                continue
            s, p, q = D["S"][o, ii].sum(), D["P"][o, ii].sum(), D["Q"][o, ii].sum()
            share = float(np.clip(O.api_share[o] * np.exp(rng.normal(0.0, 0.05)), 0.0, 0.95))
            for ch, f in (("app", 1.0 - share), ("api", share)):
                rr = int(round(r * f))
                if rr == 0:
                    continue
                out.append([d.isoformat(), O.ids[o], O.names[o], "docs", ch, rr, int(round(s * f)), int(round(p * f)),
                            int(round(q * f))])
        pr = W.pilot["calls"][ii].sum()
        if pr > 0:
            o = W.pilot_org
            out.append([d.isoformat(), O.ids[o], O.names[o], "docs", "report", int(round(pr)), int(round(pr)),
                        int(round(W.pilot["prompt"][ii].sum())), int(round(W.pilot["completion"][ii].sum()))])
    out.sort(key=lambda r: (r[0], r[1], r[4]))
    return out


def node_rows(NODES):
    out = []
    for n in NODES:
        for pool, start, end in n["memberships"]:
            out.append([n["node"], n["hall"], n["rack"], "NVIDIA H100 80GB HBM3", 8, pool, start.date().isoformat(),
                        end.date().isoformat() if end else ""])
    out.sort(key=lambda r: (r[0], r[6]))
    return out


def job_objs(M):
    for e in M["jobs"]:
        yield {"ts": K.iso(e["ts"]), "job": e["job"], "node": e["node"], "pool": e["pool"], "step": e["step"],
               "result": e["result"], "detail": e["detail"]}


def console_rows(M):
    for c in M["console"]:
        yield [c["local"].strftime("%Y-%m-%d %H:%M:%S"), c["user"], c["action"], c["obj"], c["detail"], c["ticket"]]


def burn_rows(A):
    rows = sorted(A.burn, key=lambda b: (b["start"], b["node"], b["slot"], b["test"]))
    for b in rows:
        yield [b["rack"], b["node"], b["slot"], b["serial"], b["run"], b["test"], K.iso(b["start"]), K.iso(b["end"]),
               b["result"], b["note"]]


def fabric_rows(A):
    rows = sorted(A.fabric, key=lambda f: (f["start"], f["node"]))
    for f in rows:
        yield [f["rack"], f["node"], f["run"], "all_reduce_perf -b 8G -e 8G -g 8", K.iso(f["start"]), K.iso(f["end"]),
               "%.1f" % f["algbw"], "%.1f" % f["busbw"]]


def handover_rows(A):
    pos = {}
    for x in A.inventory:
        pos.setdefault(x["rack"], [])
        if x["node"] not in pos[x["rack"]]:
            pos[x["rack"]].append(x["node"])
    for x in A.inventory:
        u = 3 + 4 * pos[x["rack"]].index(x["node"])
        yield [x["rack"], "U%02d" % u, x["node"], x["slot"], x["serial"]]


def incident_rows(rows):
    for r in rows:
        yield [r["id"], K.iso(r["opened"]), K.iso(r["resolved"]), r["sev"], r["pools"], r["summary"]]


def spec_rows():
    T = K.MODEL_TAG
    return [
        ["chat-prod", T["chat-prod"], "H100 SXM5 80GB", 8, 8, 176, 256, 32768, "fp8", "vllm 0.9.2"],
        ["code-prod", T["code-prod"], "H100 SXM5 80GB", 8, 4, 160, 384, 65536, "fp8", "vllm 0.9.2"],
        ["agents-prod", T["agents-prod"], "H100 SXM5 80GB", 8, 8, 58, 192, 131072, "fp8", "vllm 0.9.2"],
        ["docs-prod", K.LC2_TAG, "H100 SXM5 80GB", 8, 8, 74, 128, 131072, "fp8", "vllm 0.9.2"],
        ["batch-prod", K.LC2_TAG, "H100 SXM5 80GB", 8, 8, 40, 192, 131072, "fp8", "vllm 0.9.2"],
        ["voice-prod", T["voice-prod"], "H100 SXM5 80GB", 8, 2, 132, 512, 16384, "fp8", "tensorrt_llm 0.20"],
        ["nx7 stopgap (leased)", "", "H100 SXM5 80GB", 8, "", "", "", "", "", ""],
    ]


# ------------------------------------------------------------------------------------------------- workbooks
HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
BOLD = Font(bold=True)


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_intake(path, W):
    intake = WD.q1_intake()
    wb = Workbook()
    ws = wb.active
    ws.title = "Intake Q1 2027"
    ws.append(["Fleet capacity intake, Q1 2027 (ISO weeks 2027-W01 to 2027-W13)"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Each product team files the traffic it expects in each week's busiest hour. Due 27 November 2026."])
    ws.append([])
    head = ["Pool", "Product", "Filed by", "Measure"] + ["W%02d" % i for i in range(1, 14)] + ["Filed", "Note"]
    ws.append(head)
    for c in ws[4]:
        c.font = BOLD
        c.fill = HEAD_FILL
    who = {"chat-prod": "Chat platform", "code-prod": "Code assistant", "agents-prod": "Agents platform",
           "docs-prod": "Docs", "batch-prod": "Batch API", "voice-prod": "Voice"}
    filed = {"chat-prod": "2026-11-24", "code-prod": "2026-11-26", "agents-prod": "2026-11-25",
             "docs-prod": "2026-11-27", "batch-prod": "2026-11-20", "voice-prod": "2026-11-26"}
    note = {
        "chat-prod": "Consumer app and web. January seasonal lift included.",
        "code-prod": "IDE plugin traffic including agent mode, GA on 18 January (W03).",
        "agents-prod": "Agent runs, every tool step a request.",
        "docs-prod": "App and API traffic. Report service contracts are held in the commitments ledger and are not in "
                     "these figures.",
        "batch-prod": "Asynchronous summarisation; busiest hour of the queue drain.",
        "voice-prod": "Filed as requests per second in the busiest hour.",
    }
    for p in K.POOLS:
        meas = ("Busiest-hour requests per second" if p == "voice-prod" else "Busiest-hour requests (thousands)")
        ws.append([p, K.PRODUCT[p], who[p], meas] + list(intake[p]) + [filed[p], note[p]])
    for row in ws.iter_rows(min_row=5, max_row=10, min_col=5, max_col=17):
        for c in row:
            c.number_format = "#,##0.0" if isinstance(c.value, float) and c.value != int(c.value) else "#,##0"
    _widths(ws, [13, 11, 16, 34] + [9] * 13 + [11, 70])
    ws.freeze_panes = "E5"
    wb.save(path)


def write_plan(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Allocation"
    ws.append(["Fleet plan, Q1 2027, issue v3 (1 December 2026). GPUs allocated to each pool from the Monday of each "
               "ISO week."])
    ws["A1"].font = Font(bold=True)
    ws.append([])
    head = ["Pool", "Product"] + ["W%02d (%s)" % (i + 1, m.strftime("%d %b")) for i, m in enumerate(K.Q1_WEEKS)]
    ws.append(head)
    for c in ws[3]:
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True)
    for p in K.POOLS:
        ws.append([p, K.PRODUCT[p]] + K.PLAN_Q1[p])
    tot = ["Fleet", ""] + [sum(K.PLAN_Q1[p][i] for p in K.POOLS) for i in range(13)]
    ws.append(tot)
    for c in ws[10]:
        c.font = BOLD
    _widths(ws, [13, 11] + [10] * 13)
    ws.freeze_panes = "C4"
    m = wb.create_sheet("Moves")
    m.append(["Move", "From pool", "To pool", "Nodes", "GPUs", "Effective", "Reason"])
    for c in m[1]:
        c.font = BOLD
        c.fill = HEAD_FILL
    m.append(["MV-2027-02", "batch-prod", "docs-prod", 6, 48, "2027-02-15", "Docs growth; Batch queue holds its "
              "drain time at 280 GPUs"])
    _widths(m, [12, 12, 12, 7, 7, 11, 60])
    a = wb.create_sheet("About")
    for line in [
        ["Issue", "v3"], ["Issued", "2026-12-01"], ["Owner", K.PEOPLE["capacity_lead"] + ", fleet capacity"],
        ["Covers", "ISO weeks 2027-W01 (Monday 4 January) to 2027-W13 (Sunday 4 April)"],
        ["Not in this plan", "The leased stopgap cluster (240 GPUs) handed over on 4 January 2027."],
        ["Supersedes", "v2 (13 November 2026): Docs from 15 February at 656, not 640."],
    ]:
        a.append(line)
    _widths(a, [18, 90])
    wb.save(path)


def write_ledger(path, W):
    wb = Workbook()
    ws = wb.active
    ws.title = "Contracts"
    head = ["Contract", "Account", "Account id", "Product", "Service", "Tier", "Contracted peak (reports/hour)",
            "Q1 committed reports", "Q1 calls", "Signed", "Go-live", "Status", "Account owner"]
    ws.append(head)
    for c in ws[1]:
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    owners = ["Demi Verhagen", "Vincent Palman", "Sam Kolen"]
    for k_, c in enumerate(K_ledger(W)):
        ws.append([c["ref"], W.orgs.names[c["acc"]], W.orgs.ids[c["acc"]], "Docs", "Report service", c["tier"],
                   c["peak"], c["q1_reports"] if c["q1_reports"] else None, c["q1_calls"] if c["q1_calls"] else None,
                   c["signed"].isoformat() if c["signed"] else "", c["golive"].isoformat(), c["status"],
                   owners[k_ % 3]])
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=9):
        for cell in row:
            cell.number_format = "#,##0"
    _widths(ws, [13, 32, 14, 8, 15, 7, 13, 13, 12, 11, 11, 18, 15])
    ws.freeze_panes = "B2"
    t = wb.create_sheet("Tiers")
    t.append(["Tier", "Calls per report", "Prompt tokens per call (limit)", "Completion tokens per call (limit)",
              "List price per report (EUR)", "Use"])
    for c in t[1]:
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    use = {"RS-6": "short reports: claim files, case summaries", "RM-8": "standard reports: dossiers, audit files",
           "RL-12": "long reports: due diligence, full-file reviews"}
    for tier, (n, pm, qm, price) in K.TIERS.items():
        t.append([tier, n, pm, qm, price, use[tier]])
    _widths(t, [8, 10, 16, 16, 14, 50])
    n = wb.create_sheet("Notes")
    for line in [["Owner", K.PEOPLE["reports_pm"] + ", report service"], ["Exported", "2026-12-04"],
                 ["Calls", "A report is generated section by section; each section is one call to the Docs "
                           "model. Q1 calls are committed reports times the tier's calls per report."]]:
        n.append(line)
    _widths(n, [12, 100])
    wb.save(path)


def K_ledger(W):
    return WD.contracts()
