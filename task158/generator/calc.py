"""task158 calc: every figure computed from the shipped files alone, the way an analyst would: the golden, every
rung of the ladder, both asks and their stops, and the out-of-sample check on the early-access tenant. Reads only
<target>/; never imports the world."""
from __future__ import annotations

import csv
import datetime as dt
import json
import os
from collections import defaultdict

import numpy as np
from openpyxl import load_workbook

import common as K

F_ROLLUP = "fleet_rollup_hourly_2026-01-01_2026-12-06.csv"
F_BILLING = "docs_usage_by_org_daily_2026-09-01_2026-12-06.csv"
F_INTAKE = "q1_2027_capacity_intake.xlsx"
F_PLAN = "fleet_plan_2027Q1.xlsx"
F_LEDGER = "report_service_commitments_2026-12-04.xlsx"
F_NODES = "fleet_nodes_2026-12-06.csv"
F_JOBS = "fleetctl_jobs_2026-10-26_2026-12-06.jsonl"
F_CONSOLE = "ops_console_audit_2026-10-26_2026-12-06.csv"
F_BURN = "nx7_burnin_results.csv"
F_FABRIC = "nx7_fabric_results.csv"
F_HANDOVER = "nx7_handover_inventory_2026-12-04.csv"


def _t(s):
    return dt.datetime.strptime(s, "%Y-%m-%dT%H:%M:%SZ")


# ------------------------------------------------------------------------------------------------- readers
def read_rollup(target):
    rows = defaultdict(lambda: defaultdict(list))
    with open(os.path.join(target, F_ROLLUP), encoding="utf-8") as f:
        for r in csv.DictReader(f):
            p = r["pool"]
            d = rows[p]
            d["t"].append(_t(r["hour_start_utc"]))
            d["model"].append(r["model"])
            for c in ("gpus_in_service", "gpu_busy_s", "requests", "sessions_new", "prompt_tokens", "completion_tokens"):
                d[c].append(int(r[c]))
    out = {}
    for p, d in rows.items():
        out[p] = {k: (np.array(v) if k not in ("t", "model") else v) for k, v in d.items()}
    return out


def read_intake(target):
    wb = load_workbook(os.path.join(target, F_INTAKE), data_only=True)
    ws = wb["Intake Q1 2027"]
    out, unit = {}, {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] in K.POOLS:
            vals = [float(x) for x in row[4:17]]
            if "per second" in row[3]:
                vals = [v * 3600.0 for v in vals]
            else:
                vals = [v * 1000.0 for v in vals]
            out[row[0]] = vals
            unit[row[0]] = row[3]
    return out, unit


def read_plan(target):
    wb = load_workbook(os.path.join(target, F_PLAN), data_only=True)
    ws = wb["Allocation"]
    out = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0] in K.POOLS:
            out[row[0]] = [int(x) for x in row[2:15]]
    return out


def read_ledger(target):
    wb = load_workbook(os.path.join(target, F_LEDGER), data_only=True)
    ws = wb["Contracts"]
    head = [c.value for c in ws[1]]
    rows = [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    tiers = {}
    for r in wb["Tiers"].iter_rows(min_row=2, values_only=True):
        if r[0]:
            tiers[r[0]] = dict(calls=int(r[1]), pmax=int(r[2]), qmax=int(r[3]))
    return rows, tiers


def read_billing(target):
    out = []
    with open(os.path.join(target, F_BILLING), encoding="utf-8") as f:
        for r in csv.DictReader(f):
            out.append(dict(date=dt.date.fromisoformat(r["usage_date"]), org=r["org_id"], name=r["org_name"],
                            channel=r["channel"], requests=int(r["requests"]), sessions=int(r["sessions"]),
                            input=int(r["input_tokens"]), output=int(r["output_tokens"])))
    return out


def read_nodes(target):
    nodes = {}
    with open(os.path.join(target, F_NODES), encoding="utf-8") as f:
        for r in csv.DictReader(f):
            n = nodes.setdefault(r["node"], dict(node=r["node"], memberships=[]))
            start = dt.datetime.fromisoformat(r["member_from"])
            end = dt.datetime.fromisoformat(r["member_to"]) if r["member_to"] else None
            n["memberships"].append([r["pool"], start, end])
    return nodes


def read_jobs(target):
    out = []
    with open(os.path.join(target, F_JOBS), encoding="utf-8") as f:
        for line in f:
            e = json.loads(line)
            e["ts"] = _t(e["ts"])
            out.append(e)
    return out


def read_console(target):
    out = []
    with open(os.path.join(target, F_CONSOLE), encoding="utf-8") as f:
        for r in csv.DictReader(f):
            out.append(dict(local=dt.datetime.strptime(r["time (Europe/Amsterdam)"], "%Y-%m-%d %H:%M:%S"),
                            user=r["user"], action=r["action"], obj=r["object"], detail=r["detail"], ticket=r["ticket"]))
    return out


# ------------------------------------------------------------------------------------------------- the close-out
def closeout(target):
    R = read_rollup(target)
    q0, q1 = dt.datetime(2026, 7, 1), dt.datetime(2026, 10, 1)
    out = {"pools": {}}
    fleet_busy = fleet_alloc = 0.0
    for p in K.POOLS:
        d = R[p]
        m = np.array([q0 <= t < q1 for t in d["t"]])
        busy = d["gpu_busy_s"][m].sum()
        req = d["requests"][m].sum()
        hrs = [t for t, k in zip(d["t"], m) if k]
        alloc_h = sum(K.alloc_at(p, t + dt.timedelta(minutes=59, seconds=59)) for t in hrs)
        i = int(np.argmax(np.where(m, d["gpu_busy_s"], -1)))
        months = []
        for mo in (7, 8, 9):
            mm = np.array([t.month == mo and t.year == 2026 for t in d["t"]])
            months.append(d["gpu_busy_s"][mm].sum() / 3600.0)
        out["pools"][p] = dict(alloc=K.alloc_at(p, dt.datetime(2026, 9, 30, 23, 59)), busy_h=busy / 3600.0,
                               share=100.0 * (busy / 3600.0) / alloc_h, req=int(req),
                               sess=int(d["sessions_new"][m].sum()), ptok=int(d["prompt_tokens"][m].sum()),
                               ctok=int(d["completion_tokens"][m].sum()), gpr=busy / req,
                               peak_hour=d["t"][i].strftime("%d %b %H:00"), peak_gpus=int(round(d["gpu_busy_s"][i] / 3600.0)),
                               months=months, sep_peak_share=None)
        sep = np.array([t.month == 9 and t.year == 2026 for t in d["t"]])
        j = int(np.argmax(np.where(sep, d["gpu_busy_s"], -1)))
        out["pools"][p]["sep_peak_share"] = d["gpu_busy_s"][j] / 3600.0 / K.alloc_at(p, d["t"][j]) * 100.0
        fleet_busy += busy / 3600.0
        fleet_alloc += alloc_h
    out["fleet_busy_h"] = int(round(fleet_busy))
    out["fleet_share"] = 100.0 * fleet_busy / fleet_alloc
    return out


def closeout_gpr(target):
    """GPU-seconds per request as the close-out prints them (three decimals), read back from the rollup."""
    C = closeout(target)
    return {p: round(C["pools"][p]["gpr"], 3) for p in K.POOLS}, C


# ------------------------------------------------------------------------------------------------- cost of a report call
def fit_two_rates(R, pool="batch-prod", start=None, end=None, intercept=False, model=K.LC2_TAG):
    d = R[pool]
    m = np.array([(mdl == model) for mdl in d["model"]])
    if start:
        m &= np.array([t >= start for t in d["t"]])
    if end:
        m &= np.array([t < end for t in d["t"]])
    X = np.c_[d["prompt_tokens"][m], d["completion_tokens"][m]].astype(float)
    y = d["gpu_busy_s"][m].astype(float)
    if intercept:
        X = np.c_[np.ones(len(y)), X]
        c = np.linalg.lstsq(X, y, rcond=None)[0]
        return float(c[1]), float(c[2]), float(c[0])
    c = np.linalg.lstsq(X, y, rcond=None)[0]
    return float(c[0]), float(c[1]), 0.0


def fit_pooled_models(R, pool="batch-prod"):
    """The same two-rate fit over every hour of the year, lc1 and lc2 together."""
    d = R[pool]
    X = np.c_[d["prompt_tokens"], d["completion_tokens"]].astype(float)
    c = np.linalg.lstsq(X, d["gpu_busy_s"].astype(float), rcond=None)[0]
    return float(c[0]), float(c[1])


# ------------------------------------------------------------------------------------------------- the main call
def contracts_live(rows, tiers):
    out = []
    for r in rows:
        if r["Status"] != "Signed":
            continue
        t = tiers[r["Tier"]]
        out.append(dict(ref=r["Contract"], tier=r["Tier"], peak=int(r["Contracted peak (reports/hour)"]),
                        golive=dt.date.fromisoformat(r["Go-live"]), calls=t["calls"], pmax=t["pmax"], qmax=t["qmax"]))
    return out


def report_work_per_hour(live, week_monday, per_report):
    return sum(c["peak"] * per_report(c) for c in live if c["golive"] <= week_monday)


def shortfalls(intake, plan, gpr, live, per_report, docs_interactive_gpr=None, alloc_override=None):
    """Each pool's weekly shortfall and its Q1 peak-week shortfall under a costing of report calls."""
    out = {}
    for p in K.POOLS:
        w = []
        for i, monday in enumerate(K.Q1_WEEKS):
            g = docs_interactive_gpr if (p == "docs-prod" and docs_interactive_gpr) else gpr[p]
            work = intake[p][i] * g
            if p == "docs-prod":
                work += report_work_per_hour(live, monday, per_report)
            alloc = (alloc_override or plan)[p][i]
            w.append(work / K.CARRY - alloc)
        out[p] = dict(weekly=w, peak=max(w), week=1 + int(np.argmax(w)))
    pick = max(K.POOLS, key=lambda p: out[p]["peak"])
    return out, pick


def main_call(target):
    R = read_rollup(target)
    intake, unit = read_intake(target)
    plan = read_plan(target)
    rows, tiers = read_ledger(target)
    gpr, C = closeout_gpr(target)
    gpr_exact = {p: C["pools"][p]["gpr"] for p in K.POOLS}
    a, b, _ = fit_two_rates(R)
    live = contracts_live(rows, tiers)
    golden_cost = lambda c: c["calls"] * (a * c["pmax"] + b * c["qmax"])
    sf, pick = shortfalls(intake, plan, gpr_exact, live, golden_cost)
    docs = sf["docs-prod"]
    full = report_work_per_hour(live, K.Q1_WEEKS[-1], golden_cost) / K.CARRY
    inter13 = intake["docs-prod"][12] * gpr_exact["docs-prod"] / K.CARRY
    return dict(R=R, intake=intake, unit=unit, plan=plan, rows=rows, tiers=tiers, gpr=gpr, gpr_exact=gpr_exact, C=C,
                a=a, b=b, live=live, sf=sf, pick=pick, docs_peak=docs["peak"], docs_week=docs["week"],
                reports_full=full, docs_inter_w13=inter13, cost=golden_cost)


# ------------------------------------------------------------------------------------------------- the rungs
def rungs(M, target):
    """Every route on the ladder, each costing report calls its own way on the same literal estimand."""
    R, intake, plan, live = M["R"], M["intake"], M["plan"], M["live"]
    gpr = M["gpr_exact"]
    d = R["docs-prod"]
    q3 = np.array([dt.datetime(2026, 7, 1) <= t < dt.datetime(2026, 10, 1) for t in d["t"]])
    busy, req = d["gpu_busy_s"][q3].sum(), d["requests"][q3].sum()
    ptok, ctok, sess = d["prompt_tokens"][q3].sum(), d["completion_tokens"][q3].sum(), d["sessions_new"][q3].sum()
    per_tok = busy / (ptok + ctok)
    X = np.c_[d["prompt_tokens"][q3], d["completion_tokens"][q3]].astype(float)
    r2a, r2b = np.linalg.lstsq(X, d["gpu_busy_s"][q3].astype(float), rcond=None)[0]
    X3 = np.c_[d["sessions_new"][q3], d["requests"][q3] - d["sessions_new"][q3], d["completion_tokens"][q3]].astype(float)
    ns, nf, nb = np.linalg.lstsq(X3, d["gpu_busy_s"][q3].astype(float), rcond=None)[0]
    per_session = busy / sess
    bq = R["batch-prod"]
    bm = np.array([dt.datetime(2026, 7, 1) <= t < dt.datetime(2026, 10, 1) for t in bq["t"]])
    batch_req = bq["gpu_busy_s"][bm].sum() / bq["requests"][bm].sum()
    batch_tok = bq["gpu_busy_s"][bm].sum() / (bq["prompt_tokens"][bm].sum() + bq["completion_tokens"][bm].sum())
    P = pilot(M, target)
    a, b = M["a"], M["b"]
    pa, pb = fit_pooled_models(R)
    brief = 260.0
    routes = {
        "R0 Docs GPU-seconds per request": lambda c: c["calls"] * gpr["docs-prod"],
        "R1a Docs blended rate per token": lambda c: c["calls"] * (c["pmax"] + c["qmax"]) * per_tok,
        "R1b two rates fitted on Docs' raw tokens": lambda c: c["calls"] * (r2a * c["pmax"] + r2b * c["qmax"]),
        "R2a Batch GPU-seconds per request": lambda c: c["calls"] * batch_req,
        "R2b per new session (Docs' session-start cost)": lambda c: c["calls"] * (ns + nb * c["qmax"]),
        "R2c per session (Docs' GPU-seconds per session)": lambda c: c["calls"] * per_session,
        "R3a the tenant's cost per report": lambda c: P["per_report"],
        "R3b the tenant's cost per call": lambda c: c["calls"] * P["per_call"],
        "R3c the pack shared across a report's calls": lambda c: (a * c["pmax"] + (c["calls"] - 1) * a * brief
                                                                  + c["calls"] * b * c["qmax"]),
        "R4 computed prompt and completion tokens at lc2's rates": M["cost"],
        "R5a Batch blended rate per token": lambda c: c["calls"] * (c["pmax"] + c["qmax"]) * batch_tok,
        "R5b the tenant's blended rate per token": lambda c: c["calls"] * (c["pmax"] + c["qmax"]) * P["per_token"],
        "R5c lc1 and lc2 hours pooled in the Batch fit": lambda c: c["calls"] * (pa * c["pmax"] + pb * c["qmax"]),
    }
    out = {}
    for name, f in routes.items():
        sf, pick = shortfalls(intake, plan, gpr, live, f)
        out[name] = dict(pick=pick, docs=sf["docs-prod"]["peak"], code=sf["code-prod"]["peak"],
                         reports=report_work_per_hour(live, K.Q1_WEEKS[-1], f) / K.CARRY,
                         ranks=sorted(K.POOLS, key=lambda p: -sf[p]["peak"]))
    return out


# ------------------------------------------------------------------------------------------------- the early-access tenant
def pilot(M, target):
    """The early-access tenant's report calls from billing, set against Docs' busy time on the tenant's days: the
    excess over Docs' own interactive work (its Q3 GPU-seconds per request) against each costing's prediction."""
    bill = read_billing(target)
    rows, _ = read_ledger(target)
    ea = [r for r in rows if r["Contract"].startswith("EA")][0]
    org = ea["Account id"]
    rep = [x for x in bill if x["org"] == org and x["channel"] == "report"]
    by_day = {x["date"]: x for x in rep}
    R = M["R"]
    d = R["docs-prod"]
    day_of = [K.utc_to_ams(t).date() for t in d["t"]]
    busy_d, req_d = defaultdict(float), defaultdict(float)
    for t, dd, bsy, rq in zip(d["t"], day_of, d["gpu_busy_s"], d["requests"]):
        busy_d[dd] += bsy
        req_d[dd] += rq
    g = M["gpr_exact"]["docs-prod"]
    days = sorted(dd for dd in by_day if dt.date(2026, 10, 5) <= dd <= dt.date(2026, 12, 4))
    calls = sum(by_day[dd]["requests"] for dd in days)
    ptk = sum(by_day[dd]["input"] for dd in days)
    ctk = sum(by_day[dd]["output"] for dd in days)
    excess = sum(busy_d[dd] - g * (req_d[dd] - by_day[dd]["requests"]) for dd in days)
    a, b = M["a"], M["b"]
    preds = {"computed prompt and completion at lc2's rates": a * ptk + b * ctk,
             "Docs GPU-seconds per request": g * calls,
             "the pack computed once a report": a * ptk / 6.0 + a * 260.0 * calls * 5.0 / 6.0 + b * ctk}
    reports = calls / 6.0
    out = dict(days=len(days), calls=calls, ptk=ptk, ctk=ctk, excess=excess, preds=preds,
               ratio={k: v / excess for k, v in preds.items()},
               per_call=excess / calls, per_report=excess / reports, per_token=excess / (ptk + ctk))
    # a two-rate fit on the tenant's own days: its prompt and completion move together
    y = np.array([busy_d[dd] - g * (req_d[dd] - by_day[dd]["requests"]) for dd in days])
    X = np.array([[by_day[dd]["input"], by_day[dd]["output"]] for dd in days], dtype=float)
    coef = np.linalg.lstsq(X, y, rcond=None)[0]
    out["own_fit"] = (float(coef[0]), float(coef[1]))
    out["corr_daily_tokens"] = float(np.corrcoef(X[:, 0] / np.maximum(X[:, 0].mean(), 1), X[:, 1])[0, 1])
    return out


def pilot_ratio(M, target, utc_days=False, gpr_docs=None):
    """lc2's prediction of the tenant's excess busy time over the excess itself, and Docs' rate per request over the
    excess: billing's Amsterdam days matched to the rollup's Amsterdam days (or to UTC days), Docs' interactive rate
    exact (or as the close-out prints it)."""
    bill = read_billing(target)
    rows, _ = read_ledger(target)
    org = [r for r in rows if r["Contract"].startswith("EA")][0]["Account id"]
    rep = {x["date"]: x for x in bill if x["org"] == org and x["channel"] == "report"}
    d = M["R"]["docs-prod"]
    busy, req = defaultdict(float), defaultdict(float)
    for t, b, r in zip(d["t"], d["gpu_busy_s"], d["requests"]):
        day = t.date() if utc_days else K.utc_to_ams(t).date()
        busy[day] += b
        req[day] += r
    g = gpr_docs or M["gpr_exact"]["docs-prod"]
    days = sorted(rep)
    exc = sum(busy[dd] - g * (req[dd] - rep[dd]["requests"]) for dd in days)
    pred = sum(M["a"] * rep[dd]["input"] + M["b"] * rep[dd]["output"] for dd in days)
    return pred / exc, g * sum(rep[dd]["requests"] for dd in days) / exc


# ------------------------------------------------------------------------------------------------- Ask A
def _pool_of(node, t):
    for pool, start, end in node["memberships"]:
        if t >= start and (end is None or t < end):
            return pool
    return None


def ask_a(target, console_as_utc=False, use_console=True, use_job_complete=True, current_pool=False,
          drop_unpaired=False):
    nodes = read_nodes(target)
    jobs = read_jobs(target)
    console = read_console(target) if use_console else []
    job_node = {e["job"]: e["node"] for e in jobs}
    ev = []
    for e in jobs:
        if e["step"] == "cordon" and e["result"] == "ok":
            ev.append((e["ts"], e["node"], 1))
        elif e["step"] == "uncordon" and e["result"] == "ok":
            ev.append((e["ts"], e["node"], 0))
    for c in console:
        t = c["local"] if console_as_utc else K.ams_to_utc(c["local"])
        if c["action"] == "node.cordon":
            ev.append((t, c["obj"], 1))
        elif c["action"] == "node.uncordon":
            ev.append((t, c["obj"], 0))
        elif c["action"] == "fleetctl.job.complete" and use_job_complete:
            ev.append((t, job_node[c["obj"]], 0))
    ev.sort()
    state, stays = {}, []
    for t, node, kind in ev:
        if kind == 1 and node not in state:
            state[node] = t
        elif kind == 0 and node in state:
            stays.append((node, state.pop(node), t))
    if not drop_unpaired:
        for node, t in state.items():
            stays.append((node, t, K.MAINT_WINDOW[1]))
    gh = {p: 0.0 for p in K.POOLS}
    longest = {p: 0.0 for p in K.POOLS}
    for node, s, e in stays:
        n = nodes[node]
        a, b = max(s, K.NOV[0]), min(e, K.NOV[1])
        if b <= a:
            continue
        if current_pool:
            p = [m[0] for m in n["memberships"] if m[2] is None][0]
            gh[p] += 8 * (b - a).total_seconds() / 3600.0
        else:
            cuts = sorted({a, b} | {m[1] for m in n["memberships"] if a < m[1] < b})
            for x, y in zip(cuts[:-1], cuts[1:]):
                gh[_pool_of(n, x)] += 8 * (y - x).total_seconds() / 3600.0
        if s >= K.NOV[0] and e <= K.NOV[1]:
            p = ([m[0] for m in n["memberships"] if m[2] is None][0] if current_pool else _pool_of(n, s))
            longest[p] = max(longest[p], (e - s).total_seconds() / 3600.0)
    return gh, longest


def ask_a_rollup(target):
    """The rollup's end-of-hour samples read as the out-of-service record."""
    R = read_rollup(target)
    plan_q4 = {p: [] for p in K.POOLS}
    gh = {}
    for p in K.POOLS:
        d = R[p]
        tot = 0.0
        for t, g in zip(d["t"], d["gpus_in_service"]):
            if K.NOV[0] <= t < K.NOV[1]:
                tot += K.alloc_at(p, t + dt.timedelta(minutes=59, seconds=59)) - g
        gh[p] = tot
    return gh


# ------------------------------------------------------------------------------------------------- Ask B
def ask_b(target, gbit_fix=True, by_serial=True, racks_from_results=False, divide_all=False):
    burn = list(csv.DictReader(open(os.path.join(target, F_BURN), encoding="utf-8")))
    fab = list(csv.DictReader(open(os.path.join(target, F_FABRIC), encoding="utf-8")))
    inv = list(csv.DictReader(open(os.path.join(target, F_HANDOVER), encoding="utf-8")))
    review_end = dt.datetime(2026, 12, 4, 0, 0)
    fitted = defaultdict(dict)
    hand_rack = {}
    for x in inv:
        fitted[x["node"]][int(x["slot"])] = x["gpu_serial"]
        hand_rack[x["node"]] = x["rack"]
    # burn-in: the last run of each test on each GPU serial
    last = {}
    for b in burn:
        if _t(b["finished_utc"]) >= review_end:
            continue
        key = (b["gpu_serial"], b["test"]) if by_serial else (b["node"], int(b["slot"]), b["test"])
        prev = last.get(key)
        if prev is None or int(b["run"]) > int(prev["run"]):
            last[key] = b
    gpu_ok = {}
    for node, slots in fitted.items():
        ok = True
        for slot, ser in slots.items():
            for test in ("dcgm_diag_r3", "gemm_soak_4h", "hbm_memtest"):
                key = (ser, test) if by_serial else (node, slot, test)
                r = last.get(key)
                if r is None or r["result"] != "PASS":
                    ok = False
        gpu_ok[node] = ok
    fab_last = {}
    for f in fab:
        if _t(f["finished_utc"]) >= review_end:
            continue
        prev = fab_last.get(f["node"])
        if prev is None or int(f["run"]) > int(prev["run"]):
            fab_last[f["node"]] = f
    accepted, held = defaultdict(int), []
    result_rack = {}
    for b in burn:
        result_rack.setdefault(b["node"], b["rack"])
    for node in fitted:
        f = fab_last.get(node)
        bw = None
        if f is not None:
            bw = float(f["busbw"])
            if divide_all:
                bw = bw / 8.0
            elif gbit_fix and _t(f["started_utc"]) >= K.HARNESS_CHANGE:
                bw = bw / 8.0
        ok = gpu_ok[node] and bw is not None and bw >= 360.0
        rack = result_rack[node] if racks_from_results else hand_rack[node]
        if ok:
            accepted[rack] += 8
        else:
            held.append(node)
    for r in ("R11", "R12", "R13", "R14"):
        accepted.setdefault(r, 0)
    return dict(sorted(accepted.items())), sorted(held)
