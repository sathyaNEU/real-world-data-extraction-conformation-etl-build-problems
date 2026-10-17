"""task158 verify: an independent recomputation from the shipped bundle alone, on its own code path (pandas, no
generator module imported): the call and its critical components, the default path and three other rungs, the
out-of-sample check on the early-access tenant, Ask A and Ask B, each set against generator/figures.json and the golden
workbook.

    python3 task158/generator/verify.py [target_dir]"""
from __future__ import annotations

import json
import os
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
FAILS = []

POOLS = ["chat-prod", "code-prod", "agents-prod", "docs-prod", "batch-prod", "voice-prod"]
CARRY = 2520.0                                    # policy section 2: busy GPU-seconds of the hour divided by 2,520
MONDAYS = pd.date_range("2027-01-04", periods=13, freq="7D")
NOV0, NOV1 = pd.Timestamp("2026-11-01"), pd.Timestamp("2026-12-01")
EXPORT_END = pd.Timestamp("2026-12-07")           # the maintenance exports run to 6 December inclusive


def ok(label, cond, detail=""):
    print("%s  %-72s %s" % ("PASS" if cond else "FAIL", label, detail))
    if not cond:
        FAILS.append(label)


def to10(x):
    return int(np.floor(x / 10.0 + 0.5) * 10)


def utc_naive(s):
    return pd.to_datetime(s, utc=True).dt.tz_convert(None)


# ------------------------------------------------------------------------------------------------- the main call
def load(target):
    T = lambda f: os.path.join(target, f)
    roll = pd.read_csv(T("fleet_rollup_hourly_2026-01-01_2026-12-06.csv"))
    roll["t"] = utc_naive(roll["hour_start_utc"])
    intake = pd.read_excel(T("q1_2027_capacity_intake.xlsx"), sheet_name="Intake Q1 2027", header=3)
    intake = intake[intake["Pool"].isin(POOLS)].set_index("Pool")
    wk = ["W%02d" % i for i in range(1, 14)]
    scale = np.where(intake["Measure"].str.contains("per second"), 3600.0, 1000.0)
    req = intake[wk].astype(float).mul(scale, axis=0)
    plan = pd.read_excel(T("fleet_plan_2027Q1.xlsx"), sheet_name="Allocation", header=2)
    plan = plan[plan["Pool"].isin(POOLS)].set_index("Pool").iloc[:, 1:14].astype(int)
    plan.columns = wk
    con = pd.read_excel(T("report_service_commitments_2026-12-04.xlsx"), sheet_name="Contracts")
    tiers = pd.read_excel(T("report_service_commitments_2026-12-04.xlsx"), sheet_name="Tiers").set_index("Tier")
    return roll, req, plan, con, tiers


def q3_rate(roll):
    q3 = roll[(roll["t"] >= "2026-07-01") & (roll["t"] < "2026-10-01")]
    g = q3.groupby("pool")[["gpu_busy_s", "requests", "prompt_tokens", "completion_tokens", "sessions_new"]].sum()
    return g


def lc2_rates(roll, exact_tag=True):
    b = roll[roll["pool"] == "batch-prod"]
    b = b[b["model"] == "lc2-2026.05"] if exact_tag else b
    X = b[["prompt_tokens", "completion_tokens"]].to_numpy(float)
    y = b["gpu_busy_s"].to_numpy(float)
    return np.linalg.solve(X.T @ X, X.T @ y)


def shortfall_table(req, plan, rate, con, tiers, report_cost):
    live = con[con["Status"] == "Signed"].copy()
    live["golive"] = pd.to_datetime(live["Go-live"])
    live = live.join(tiers, on="Tier")
    live["per_report"] = [report_cost(r) for _, r in live.iterrows()]
    out = {}
    for p in POOLS:
        w = []
        for i, monday in enumerate(MONDAYS):
            busy = req.loc[p].iloc[i] * rate[p]
            if p == "docs-prod":
                on = live[live["golive"] <= monday]
                busy += float((on["Contracted peak (reports/hour)"] * on["per_report"]).sum())
            w.append(busy / CARRY - plan.loc[p].iloc[i])
        out[p] = w
    peaks = {p: max(w) for p, w in out.items()}
    pick = max(peaks, key=peaks.get)
    full = float((live["Contracted peak (reports/hour)"] * live["per_report"]).sum()) / CARRY
    return out, peaks, pick, full


def main(target):
    G = json.load(open(os.path.join(HERE, "figures.json"), encoding="utf-8"))
    roll, req, plan, con, tiers = load(target)
    g3 = q3_rate(roll)
    rate = (g3["gpu_busy_s"] / g3["requests"]).to_dict()
    a, b = lc2_rates(roll)
    calls, pmax, qmax = "Calls per report", "Prompt tokens per call (limit)", "Completion tokens per call (limit)"
    gold_cost = lambda r: r[calls] * (a * r[pmax] + b * r[qmax])
    weekly, peaks, pick, full = shortfall_table(req, plan, rate, con, tiers, gold_cost)
    wk13 = int(np.argmax(weekly[pick])) + 1
    ok("V1 the call: the pool", pick == G["call"]["pool"], pick)
    ok("V2 the call: the Q1 peak-week shortfall to the nearest 10", str(to10(peaks[pick])) == G["call"]["shortfall"],
       "%.2f" % peaks[pick])
    ok("V3 the call: the peak week", "W%02d" % wk13 == G["call"]["week"], "W%02d" % wk13)
    ok("V4 component: lc2 per computed prompt token, ms", "%.3f" % (a * 1000) == G["components"]["a_ms"],
       "%.4f" % (a * 1000))
    ok("V5 component: lc2 per completion token, ms", "%.2f" % (b * 1000) == G["components"]["b_ms"], "%.3f" % (b * 1000))
    ok("V6 component: the eleven contracts at contracted peak, GPUs", str(to10(full)) == G["components"]["reports_full"],
       "%.2f" % full)
    ok("V7 component: docs-prod's W13 allocation", str(int(plan.loc["docs-prod", "W13"])) == G["components"]["alloc_w13"])
    ok("V8 the step: Docs' Q3 GPU-seconds per request", "%.3f" % rate["docs-prod"] == G["steps"]["docs_gpr"])
    others = sorted(((v, p) for p, v in peaks.items() if p != pick), reverse=True)
    ok("V9 the runner-up and its shortfall, more than 30 GPUs behind", others[0][1] == G["call"]["runner_up"] and
       str(to10(others[0][0])) == G["call"]["runner_up_shortfall"] and peaks[pick] - others[0][0] > 30,
       "%s %.1f" % (others[0][1], others[0][0]))

    # the default path and three more rungs, recomputed here
    d3 = g3.loc["docs-prod"]
    r0 = shortfall_table(req, plan, rate, con, tiers, lambda r: r[calls] * rate["docs-prod"])
    ok("V10 the default (Docs' rate per request) files code-prod at 300",
       r0[2] == "code-prod" and to10(r0[1]["code-prod"]) == 300 and to10(r0[1]["code-prod"]) != to10(peaks[pick]),
       "Docs %.1f, Code %.1f" % (r0[1]["docs-prod"], r0[1]["code-prod"]))
    per_tok = d3["gpu_busy_s"] / (d3["prompt_tokens"] + d3["completion_tokens"])
    r1 = shortfall_table(req, plan, rate, con, tiers, lambda r: r[calls] * (r[pmax] + r[qmax]) * per_tok)
    ok("V11 Docs' blended rate per token files code-prod", r1[2] == "code-prod", "Docs %.1f" % r1[1]["docs-prod"])
    per_sess = d3["gpu_busy_s"] / d3["sessions_new"]
    r2 = shortfall_table(req, plan, rate, con, tiers, lambda r: r[calls] * per_sess)
    ok("V12 Docs' GPU-seconds per session files code-prod", r2[2] == "code-prod", "Docs %.1f" % r2[1]["docs-prod"])
    pa, pb = lc2_rates(roll, exact_tag=False)
    r5 = shortfall_table(req, plan, rate, con, tiers, lambda r: r[calls] * (pa * r[pmax] + pb * r[qmax]))
    ok("V13 lc1 and lc2 hours pooled file docs-prod in another bin", r5[2] == "docs-prod" and
       to10(r5[1]["docs-prod"]) != to10(peaks[pick]), "%.1f" % r5[1]["docs-prod"])

    # the out-of-sample check on the early-access tenant
    bill = pd.read_csv(os.path.join(target, "docs_usage_by_org_daily_2026-09-01_2026-12-06.csv"))
    ea = con[con["Contract"].str.startswith("EA")].iloc[0]
    ten = bill[(bill["org_id"] == ea["Account id"]) & (bill["channel"] == "report")].copy()
    ten["day"] = pd.to_datetime(ten["usage_date"]).dt.date
    docs = roll[roll["pool"] == "docs-prod"].copy()
    docs["day"] = docs["t"].dt.tz_localize("UTC").dt.tz_convert("Europe/Amsterdam").dt.date
    dd = docs.groupby("day")[["gpu_busy_s", "requests"]].sum()
    tday = ten.set_index("day")
    other = dd.loc[tday.index, "requests"] - tday["requests"]          # Docs' own requests on the tenant's days
    excess = float((dd.loc[tday.index, "gpu_busy_s"] - rate["docs-prod"] * other).sum())
    pred = float(a * ten["input_tokens"].sum() + b * ten["output_tokens"].sum())
    per_req = float(rate["docs-prod"] * ten["requests"].sum())
    ok("V14 lc2's rates predict the tenant's excess busy time within the write-up's bound",
       abs(pred / excess - 1) * 100 <= int(G["steps"]["pilot_within_pct"]), "%.4f" % (pred / excess))
    ok("V15 Docs' rate per request predicts under a third of it", per_req / excess < 0.33, "%.3f" % (per_req / excess))
    ok("V16 the tenant's report calls in billing", "{:,}".format(int(ten["requests"].sum())) == G["steps"]["pilot_calls"])
    ok("V17 every tenant report call is its own session", bool((ten["requests"] == ten["sessions"]).all()))

    # ------------------------------------------------------------------------------------------- Ask A
    nodes = pd.read_csv(os.path.join(target, "fleet_nodes_2026-12-06.csv"))
    nodes["start"] = pd.to_datetime(nodes["member_from"])
    nodes["end"] = pd.to_datetime(nodes["member_to"]).fillna(pd.Timestamp("2100-01-01"))
    jobs = pd.read_json(os.path.join(target, "fleetctl_jobs_2026-10-26_2026-12-06.jsonl"), lines=True,
                        convert_dates=False)
    jobs["t"] = utc_naive(jobs["ts"])
    cons = pd.read_csv(os.path.join(target, "ops_console_audit_2026-10-26_2026-12-06.csv"))
    cons["t"] = pd.to_datetime(cons["time (Europe/Amsterdam)"]).dt.tz_localize("Europe/Amsterdam").dt.tz_convert(
        "UTC").dt.tz_convert(None)
    node_of_job = jobs.drop_duplicates("job").set_index("job")["node"]
    ev = [jobs[(jobs["step"] == "cordon") & (jobs["result"] == "ok")].assign(kind="open")[["t", "node", "kind"]],
          jobs[(jobs["step"] == "uncordon") & (jobs["result"] == "ok")].assign(kind="close")[["t", "node", "kind"]],
          cons[cons["action"] == "node.cordon"].rename(columns={"object": "node"}).assign(kind="open")[["t", "node",
                                                                                                         "kind"]],
          cons[cons["action"] == "node.uncordon"].rename(columns={"object": "node"}).assign(kind="close")[["t", "node",
                                                                                                           "kind"]]]
    jc = cons[cons["action"] == "fleetctl.job.complete"].copy()
    jc["node"] = jc["object"].map(node_of_job)
    ev.append(jc.assign(kind="close")[["t", "node", "kind"]])
    ev = pd.concat(ev).sort_values(["node", "t", "kind"], kind="mergesort")
    stays = []
    for node, grp in ev.groupby("node", sort=True):
        start = None
        for t, kind in zip(grp["t"], grp["kind"]):
            if kind == "open" and start is None:
                start = t
            elif kind == "close" and start is not None:
                stays.append((node, start, t))
                start = None
        if start is not None:
            stays.append((node, start, EXPORT_END))
    gh = {p: 0.0 for p in POOLS}
    longest_in, longest_touch, longest_clip = ({p: 0.0 for p in POOLS} for _ in range(3))
    mem = {n: g.sort_values("start")[["pool", "start", "end"]].to_numpy() for n, g in nodes.groupby("node")}
    for node, s, e in stays:
        lo_, hi_ = max(s, NOV0), min(e, NOV1)
        if hi_ <= lo_:
            continue
        for pool, ms, me in mem[node]:
            x, y = max(lo_, ms), min(hi_, me)
            if y > x:
                gh[pool] += 8 * (y - x).total_seconds() / 3600.0
        p0 = [pool for pool, ms, me in mem[node] if ms <= max(s, NOV0) < me][0]
        longest_clip[p0] = max(longest_clip[p0], (hi_ - lo_).total_seconds() / 3600.0)
        longest_touch[p0] = max(longest_touch[p0], (e - s).total_seconds() / 3600.0)
        if s >= NOV0 and e <= NOV1:
            longest_in[p0] = max(longest_in[p0], (e - s).total_seconds() / 3600.0)
    A = G["answers"]["out_of_service"]
    ok("V18 Ask A: GPU-hours out of service, every pool", all("{:,}".format(to10(gh[p])) == A[p]["gpu_hours"]
                                                               for p in POOLS), " ".join("%.1f" % gh[p] for p in POOLS))
    ok("V19 Ask A: longest stay, every pool", all("%.1f" % longest_in[p] == A[p]["longest"] for p in POOLS),
       " ".join("%.2f" % longest_in[p] for p in POOLS))
    ok("V20 Ask A: stays crossing a month edge, whole or clipped, change no pool's longest",
       all("%.1f" % longest_touch[p] == "%.1f" % longest_in[p] == "%.1f" % longest_clip[p] for p in POOLS))

    # ------------------------------------------------------------------------------------------- Ask B
    inv = pd.read_csv(os.path.join(target, "nx7_handover_inventory_2026-12-04.csv"))
    burn = pd.read_csv(os.path.join(target, "nx7_burnin_results.csv"))
    fab = pd.read_csv(os.path.join(target, "nx7_fabric_results.csv"))
    last = burn.sort_values("run").groupby(["gpu_serial", "test"]).tail(1).set_index(["gpu_serial", "test"])["result"]
    tests = sorted(burn["test"].unique())
    fab = fab.sort_values("run").groupby("node").tail(1).set_index("node")
    gbps = np.where(fab["busbw"] > 1000, fab["busbw"] / 8.0, fab["busbw"])      # Gbit/s runs read by their size
    fab_ok = pd.Series(gbps >= 360.0, index=fab.index)
    acc, held = {r: 0 for r in ("R11", "R12", "R13", "R14")}, []
    for node, g in inv.groupby("node"):
        burn_ok = all(last.get((s, t_), "MISSING") == "PASS" for s in g["gpu_serial"] for t_ in tests)
        if burn_ok and bool(fab_ok.get(node, False)) and len(g) == 8:
            acc[g["rack"].iloc[0]] += 8
        else:
            held.append(node)
    B = G["answers"]["acceptance"]
    ok("V21 Ask B: accepted GPUs, every rack", all(str(acc[r]) == B[r]["accepted"] for r in acc), str(acc))
    ok("V22 Ask B: nodes held back, every rack", sorted(held) == sorted(n for r in B for n in B[r]["held"]),
       ", ".join(sorted(held)))
    ok("V23 Ask B: the cluster total", str(sum(acc.values())) == G["answers"]["accepted_total"])

    # ------------------------------------------------------------------------------------------- the golden workbook
    wb = load_workbook(os.path.join(TASK, "golden", "stopgap_q1_2027.xlsx"), data_only=True)
    st = wb["Stopgap Q1 2027"]
    ok("V24 the golden workbook carries the call", st["A5"].value == pick and st["C5"].value == to10(peaks[pick]))
    o = {r[0]: r[1:3] for r in wb["Out of service Nov 2026"].iter_rows(min_row=4, values_only=True) if r[0] in POOLS}
    ok("V25 the golden workbook's out-of-service table", all(o[p][0] == to10(gh[p]) and
                                                             abs(o[p][1] - round(longest_in[p], 1)) < 1e-9 for p in POOLS))
    k = {r[0]: r[1] for r in wb["Stopgap acceptance"].iter_rows(min_row=4, values_only=True)
         if r[0] in ("R11", "R12", "R13", "R14")}
    ok("V26 the golden workbook's acceptance table", k == acc)
    print("verify: %d FAIL" % len(FAILS))
    return 0 if not FAILS else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(TASK, "target")))
