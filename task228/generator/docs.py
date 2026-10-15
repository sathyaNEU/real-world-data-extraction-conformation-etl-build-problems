"""task228 documents: the data-health hold standard and the reconciliation summary (pdf), the incident register (xlsx),
the data dictionary and the platform change log (md)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess
import world as W

COMPANY = W.ORG


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=13, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=14, leading=18, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=14, spaceBefore=8,
                        spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10, textColor=colors.HexColor("#444444"))
    return dict(body=body, h1=h1, h2=h2, small=small)


def _footer(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(18 * mm, 10 * mm, text)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _pdf(path, story, title, author, footer, stamp):
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title=title, author=author, creator=author, invariant=1)
    doc.build(story, onFirstPage=_footer(footer), onLaterPages=_footer(footer))
    postprocess.scrub(path, author, stamp)


def _table(rows, widths, size=8.0):
    t = Table(rows, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", size), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", size),
                           ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6ECF2")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


# ------------------------------------------------------------------ the hold standard
def standard(path):
    S = _styles()
    b, h1, h2, sm = S["body"], S["h1"], S["h2"], S["small"]
    st = [Paragraph("DH-3 Data-health hold for experiment readouts", h1),
          Paragraph("Avisly Medier A/S, Experimentation Platform and Data-health desk. Version 1.0, approved by Mathilde "
                    "Knudsen (head of the Data-health desk) on 5 October 2026, in force from Monday 4 January 2027. "
                    "Maintained by Kaspar Gregersen.", sm), Spacer(1, 6),
          Paragraph("1. Purpose", h2),
          Paragraph("From 4 January 2027 the scorecard service holds an experiment's readout automatically when the "
                    "experiment's telemetry fails a data-health check. A held readout is not published until the Data-health "
                    "desk releases it. The service watches one check. This standard sets out the checks it can watch, when a "
                    "check holds an experiment, how the desk chooses the check the service watches, and the hold line the "
                    "standard quotes for it. It replaces the weekly manual review of experiment telemetry for experiments that "
                    "start on or after 4 January 2027.", b),
          Paragraph("2. How experiments run", h2),
          Paragraph("An experiment assigns installs of the Avisly app (iOS, Android and web) to its arms by a hash of the install "
                    "ID and the experiment key, at the design share recorded in the experiment registry; an install keeps its "
                    "arm for the whole run. Every experiment runs 14 days, from a Monday to the Sunday of the following week, "
                    "Copenhagen time, and is read out once its run has ended. Installs that have not allowed analytics send no "
                    "telemetry and do not appear in the install exports; the registry's assigned counts include them.", b),
          Paragraph("3. The checks", h2),
          Paragraph("Each check compares the treatment arm with the control arm on the experiment's full run, from the export "
                    "of the experiment's installs. A session is the app's session as the SDK records it.", b)]
    rows = [["Check", "What it measures in each arm", "Holds when"],
            ["Gap sessions", Paragraph("Sessions whose events show a gap in the client's sequence numbers, as a share of "
                                       "all sessions in the arm", sm), Paragraph("the treatment's share is above the control's", sm)],
            ["Duplicate sessions", Paragraph("Sessions in which the collector received an event ID twice, as a share of all "
                                             "sessions in the arm", sm), Paragraph("the treatment's share is above the control's", sm)],
            ["Late sessions", Paragraph("Sessions whose last batch reached the collector more than six hours after the session "
                                        "ended, as a share of all sessions in the arm", sm),
             Paragraph("the treatment's share is above the control's", sm)],
            ["Events per install", Paragraph("Events the collector received from the arm's installs, divided by the arm's "
                                             "installs", sm), Paragraph("the treatment's figure is below the control's", sm)],
            ["Treatment share", Paragraph("The treatment arm's installs as a share of the experiment's exported installs", sm),
             Paragraph("the share is below the design share in the registry", sm)]]
    st += [_table(rows, [32 * mm, 92 * mm, 50 * mm]), Spacer(1, 6),
           Paragraph("4. When a check holds an experiment", h2),
           Paragraph("A check holds an experiment when, on the experiment's full run, the difference between the treatment "
                     "arm and the control arm (for the treatment share, between the arm's share and the design share) is "
                     "significant at the two-sided 1 per cent level and lies in the check's hold direction. Each check judges "
                     "each experiment once, when its run has ended.", b),
           Paragraph("5. Choosing the check the service watches", h2),
           Paragraph("The desk judges every check on the record of 2026: the 82 experiments run from January to September 2026, "
                     "every one a completed 14-day run in the registry, each counted as an experiment with a confirmed data "
                     "incident when the incident register holds a confirmed incident against it, and as an experiment without "
                     "one otherwise. The service may hold no more than one experiment in fifty that had no confirmed data "
                     "incident, so a check is within the allowance when it would have held no more than 2 per cent of the 2026 "
                     "experiments without a confirmed data incident. Of the checks within the allowance, the service watches the "
                     "one that would have held the most 2026 experiments with a confirmed data incident; where two would have "
                     "held as many, the one that would have held fewer experiments without one.", b),
           Paragraph("6. The hold line", h2),
           Paragraph("For the check the service watches, the desk quotes a hold line: the smallest difference between the arms, "
                     "in the check's own unit (percentage points for a share), at which a standard experiment would be held. A "
                     "standard experiment from 2027 runs 14 days with 10,000 installs in each arm of its export and an even "
                     "split. The line is set from both arms of the 2026 experiments that had no confirmed data incident. Owners "
                     "of larger or smaller experiments read their own readout's hold, not the line.", b),
           Paragraph("7. Records", h2),
           Paragraph("The experiment registry; the install exports (one row an install in an experiment, written by the "
                     "scorecard service at the end of each run); the incident register, the desk's log of suspected and "
                     "confirmed data incidents, where an incident is confirmed or withdrawn, and its cause established, by the "
                     "monthly reconciliation of the app's telemetry against the content API's server-side records; the monthly reconciliation summary; and "
                     "the weekly review log. Incidents that affect every install alike, whatever its arm, are logged in the "
                     "register against the whole app; no comparison between arms can see them, and the desk follows them up "
                     "through the collector's own monitoring.", b)]
    _pdf(path, st, "DH-3 Data-health hold for experiment readouts", "Kaspar Gregersen",
         "DH-3 v1.0  |  Avisly Medier A/S  |  Experimentation Platform", "2026-10-05")


# ------------------------------------------------------------------ the reconciliation summary
def reconciliation(path, rows):
    S = _styles()
    b, h1, h2, sm = S["body"], S["h1"], S["h2"], S["small"]
    st = [Paragraph("Telemetry reconciliation, January to September 2026", h1),
          Paragraph("Data-health desk, Avisly Medier A/S. Compiled by Inge Lund, 9 October 2026.", sm), Spacer(1, 6),
          Paragraph("Each month the desk reconciles the telemetry of every experiment with a suspected data incident against "
                    "the content API's server-side record of article opens. The table gives, for the arm under suspicion, the "
                    "opens the server recorded from the arm's exported installs and the opens the telemetry carries, the "
                    "finding and, for a confirmed incident, the cause the reconciliation established. A confirmed finding is "
                    "entered in the incident register as confirmed, a finding of no incident as withdrawn.", b)]
    months = sorted({r["month"] for r in rows}, key=lambda m: rows[[x["month"] for x in rows].index(m)]["month_no"])
    for m in months:
        rs = [r for r in rows if r["month"] == m]
        rd = rs[0]["recon"]
        st.append(Paragraph("%s reconciliation, run on %d %s: runs that ended in %s" % (
            rd.strftime("%B"), rd.day, rd.strftime("%B %Y"), m), h2))
        tab = [["Experiment", "Title", "Server opens", "Telemetry opens", "Finding", "Cause established"]]
        for r in rs:
            tab.append([r["key"], Paragraph(r["title"], sm), "{:,}".format(r["server"]), "{:,}".format(r["telemetry"]),
                        r["finding"], r["cause"]])
        st.append(_table(tab, [22 * mm, 52 * mm, 22 * mm, 24 * mm, 22 * mm, 32 * mm]))
    _pdf(path, st, "Telemetry reconciliation, January to September 2026", "Inge Lund",
         "Data-health desk  |  telemetry reconciliation 2026", "2026-10-09")


# ------------------------------------------------------------------ the incident register
HEAD = Font(name="Calibri", bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="35506B")


def register(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Register"
    cols = ["incident_id", "opened", "closed", "scope", "experiment", "platform", "incident_type", "status", "status_date",
            "sessions_in_window", "gap_sessions", "notes"]
    widths = [12, 17, 17, 11, 34, 10, 20, 12, 12, 18, 13, 44]
    for i, h in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = HEAD
        c.fill = FILL
        c.alignment = Alignment(vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    for r_i, r in enumerate(rows, start=2):
        for c_i, h in enumerate(cols, start=1):
            v = r.get(h)
            c = ws.cell(row=r_i, column=c_i, value=v if v not in ("",) else None)
            c.font = Font(name="Calibri", size=11)
            if h in ("sessions_in_window", "gap_sessions") and v is not None and v != "":
                c.number_format = "#,##0"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(cols)), len(rows) + 1)
    wb.properties.creator = "Inge Lund"
    wb.properties.lastModifiedBy = "Inge Lund"
    wb.save(path)
    postprocess.scrub(path, "Inge Lund", "2026-10-09")
    postprocess.finish_office(path, "2026-01-06T08:12:00Z", "2026-10-09T14:41:00Z", COMPANY)


# ------------------------------------------------------------------ the dictionary and the change log
DICTIONARY = """# Data dictionary: experiment data-health files

Avisly Medier A/S, Data-health desk. Files pulled on Monday 12 October 2026. Times are Copenhagen local time unless a
field says UTC.

## experiment_registry_2026.json

One object per experiment started in 2026, as the experiment platform holds it.

- `key`: the experiment key.
- `title`: the experiment's title as its owner entered it.
- `team`: the product team that owns the experiment.
- `start`, `end`: the run's first and last moment, Copenhagen local time with its UTC offset.
- `design`: the share of installs the experiment assigns to each arm.
- `assigned`: installs assigned to each arm over the run, including installs that have not allowed analytics.
- `status`: completed for a run that ran its 14 days.
- `scorecard_service`: the scorecard service that read the experiment out and wrote its install export (v1 for runs
  that ended by 12 April 2026, v2 after).

## install_summary_2026q1.csv

Written by scorecard service 1. One row per install in an experiment that started from 5 January to 30 March 2026,
written when the run ended. Installs that have not allowed analytics are not exported.

- `exp_id`: the experiment key.
- `install`: the install ID, hashed.
- `variant`: the install's arm; A is the control arm and B the treatment arm.
- `os`: iOS, Android or Web.
- `n_sessions`: sessions the install started during the run.
- `sessions_seq_gap`: of those, sessions whose events show a gap in the client's sequence numbers.
- `sessions_dup_ids`: of those, sessions in which the collector received an event ID twice.
- `sessions_late`: of those, sessions whose last batch reached the collector more than six hours after the session
  ended.
- `events`: events the collector received from the install during the run.

## install_summary_2026q2.parquet and install_summary_2026q3.parquet

Written by scorecard service 2, for experiments that started from 6 April to 29 June and from 6 July to 14 September
2026. The content is service 1's under new names.

- `experiment_key`, `install_id`, `arm` (control or treatment), `platform` (ios, android or web).
- `sessions`, `gap_sessions`, `duplicate_sessions`, `late_sessions`, `events_received`: as `n_sessions`,
  `sessions_seq_gap`, `sessions_dup_ids`, `sessions_late` and `events` above.

## data_incident_register_2026.xlsx

The desk's log of suspected and confirmed data incidents in 2026. An incident first logged as preliminary gets a second
row when the monthly reconciliation confirms or withdraws it.

- `incident_id`: the incident's reference.
- `opened`: when the incident was logged. `closed`: for app-wide incidents, when the collector team closed it.
- `scope`: experiment (logged against one experiment, named by its title) or app-wide (logged against the whole app,
  one row for each platform it hit).
- `experiment`: the experiment's title, typed by the desk.
- `platform`: for app-wide rows, the platform.
- `incident_type`: the type of data incident the on-call engineer logged when the incident was opened (telemetry
  loss, duplicate delivery, late delivery or consent loss).
- `status`, `status_date`: preliminary, confirmed or withdrawn, and the date the status was set.
- `sessions_in_window`: app-wide rows, sessions that started on the row's platform between opened and closed, from
  the collector.
- `gap_sessions`: app-wide rows, gap sessions the reconciliation counted.
- `notes`: free text.

## reconciliation_summary_2026.pdf

The monthly reconciliation of suspected incidents in experiments, January to September 2026 (see its first page).

## weekly_review_log_2026.csv

The Monday review of experiment telemetry, 12 January to 5 October 2026. One row per action.

- `review_week`: the ISO week of the review.
- `logged_at_utc`: when the action was logged, UTC.
- `experiment_key`: the experiment the action concerns.
- `action`: flag (the readout is held for a telemetry check), clear (the hold is lifted) or note (a comment, nothing
  held).
- `reviewer`, `comment`.

## collector_daily_2026.csv

The collector's daily totals by platform for every install, in an experiment or not: sessions started, gap,
duplicate and late sessions, events received.

## app_releases_2026.csv

App releases for iOS and Android with the SDK version each carries and the date its rollout completed.

## booking_queue_2027q1.csv

Experiments booked to start from January to March 2027, with their planned installs (both arms, before consent) and
design treatment share.

## export_manifest.csv

The files in this pull, with row counts.
"""

CHANGELOG = """# Experiment platform change log, 2026

Experimentation Platform, Avisly Medier A/S. Entries in date order.

- **5 January.** The 2026 experiment calendar opens. Every experiment runs 14 days from a Monday.
- **11 February.** The collector's certificate rotation fails for about three hours; uploads are refused and retried.
  Logged in the incident register against the whole app.
- **2 March.** Data-health desk moves the weekly review to Monday mornings.
- **30 March.** Last experiments started under scorecard service 1.
- **13 April.** Scorecard service 2 in production. Install exports move from CSV to Parquet; columns are renamed
  and the arms are called control and treatment (service 1 called them A and B). Content unchanged.
- **18 May.** New consent dialog on web.
- **1 July.** SDK 9.3 rolls out on iOS and Android.
- **3 August.** SDK 9.4.1 batching regression; fixed in 9.4.2 on 6 August.
- **14 September.** Last experiments of the 2026 record start.
- **5 October.** DH-3, the data-health hold for experiment readouts, approved; in force from 4 January 2027.
- **12 October.** Registry, install exports, incident register and review log pulled for the Data-health desk.
"""


def markdown(target, FILES):
    import os
    with open(os.path.join(target, FILES["dictionary"]), "w", encoding="utf-8", newline="\n") as f:
        f.write(DICTIONARY)
    with open(os.path.join(target, FILES["changes"]), "w", encoding="utf-8", newline="\n") as f:
        f.write(CHANGELOG)
