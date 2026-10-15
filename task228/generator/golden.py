"""The golden deliverable: auto_hold_check_scorecard_2027.xlsx (the check and its line, every check's 2026 record, the
weekly review's record by type of confirmed incident, the app-wide incidents)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import estim as E
import postprocess
import world as W

XLSX = "auto_hold_check_scorecard_2027.xlsx"
NAVY = "1F3A5F"
PICK = "E8F1E4"
THIN = Side(style="thin", color="BFBFBF")
DIRECTION = {"gap": "treatment above control", "dup": "treatment above control", "late": "treatment above control",
             "epu": "treatment below control", "share": "below the design share"}


def _head(ws, row, labels, widths):
    for i, h in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w


def _cell(ws, r, c, v, fmt=None, bold=False, fill=None, wrap=False):
    x = ws.cell(row=r, column=c, value=v)
    x.font = Font(name="Calibri", size=11, bold=bold)
    x.border = Border(bottom=THIN)
    if fmt:
        x.number_format = fmt
    if fill:
        x.fill = PatternFill("solid", fgColor=fill)
    if wrap:
        x.alignment = Alignment(wrap_text=True, vertical="top")
    return x


def write(folder, A):
    import os
    rec = A["recs"]["delta"]
    pick, ok, cap = A["picks"]["delta"]
    runner = ok[1]
    n_inc, n_none = len(A["inc"]), len(A["none"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Hold choice"
    ws["A1"] = "Scorecard hold from 4 January 2027: hold on gap sessions at %.2f points" % round(A["line"], 2)
    ws["A1"].font = Font(name="Calibri", bold=True, size=14)
    ws["A2"] = "Data-health desk, 12 October 2026. Judged on the 82 experiments of 2026 under DH-3."
    ws["A2"].font = Font(name="Calibri", italic=True, size=10, color="555555")
    rows = [("Check the scorecard service holds experiments on", E.NAMES[pick], None),
            ("Hold line for a standard experiment (10,000 installs in each arm, 14 days), percentage points",
             round(A["line"], 2), "0.00"),
            ("2026 experiments with a confirmed data incident it would have held (of %d)" % n_inc, rec[pick]["incidents"], "0"),
            ("2026 experiments without a confirmed incident it would have held (of %d)" % n_none, rec[pick]["none"], "0"),
            ("Next best check within the allowance", "%s, %d of %d" % (E.NAMES[runner], rec[runner]["incidents"], n_inc), None)]
    _head(ws, 4, ["Item", "Value"], [78, 24])
    for i, (k, v, f) in enumerate(rows, start=5):
        _cell(ws, i, 1, k)
        _cell(ws, i, 2, v, fmt=f, bold=(i == 5))
    ws["A11"] = ("Basis: every check judged on each experiment's full run at the two-sided 1 per cent level, in its hold "
                 "direction; each session share tested with its variance taken over installs, the unit the platform "
                 "assigns to arms (delta method on per-install sums), not over sessions.")
    ws["A11"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A11:B13")
    ws.freeze_panes = "A5"
    # every check
    ws2 = wb.create_sheet("Checks")
    _head(ws2, 1, ["Check", "Holds when", "With a confirmed incident, held (of %d)" % n_inc,
                   "Without one, held (of %d)" % n_none, "Within the allowance (at most %d)" % cap], [22, 28, 22, 22, 22])
    for i, m in enumerate(E.CHECKS, start=2):
        fill = PICK if m == pick else None
        _cell(ws2, i, 1, E.NAMES[m], bold=(m == pick), fill=fill)
        _cell(ws2, i, 2, DIRECTION[m], fill=fill)
        _cell(ws2, i, 3, rec[m]["incidents"], "0", fill=fill)
        _cell(ws2, i, 4, rec[m]["none"], "0", fill=fill)
        _cell(ws2, i, 5, "yes" if rec[m]["none"] <= cap else "no", fill=fill)
    ws2.freeze_panes = "A2"
    # the weekly review
    ws3 = wb.create_sheet("Weekly review")
    _head(ws3, 1, ["Type of confirmed incident (cause the reconciliation established)", "Flagged by the weekly review",
                   "Never flagged", "Confirmed incidents"], [44, 22, 16, 20])
    for i, t in enumerate(E.TYPES, start=2):
        f, n = A["review"][t]
        _cell(ws3, i, 1, t.capitalize())
        _cell(ws3, i, 2, f, "0")
        _cell(ws3, i, 3, n, "0")
        _cell(ws3, i, 4, f + n, "0")
    r = len(E.TYPES) + 2
    _cell(ws3, r, 1, "All types", bold=True)
    for j in range(3):
        _cell(ws3, r, 2 + j, sum((A["review"][t][0], A["review"][t][1], sum(A["review"][t]))[j] for t in E.TYPES), "0", bold=True)
    ws3.freeze_panes = "A2"
    # the app-wide incidents
    ws4 = wb.create_sheet("App-wide incidents")
    _head(ws4, 1, ["Incident", "Opened (Copenhagen)", "Platforms hit", "Gap sessions", "Sessions in the window",
                   "Share of the window's sessions (%)"], [12, 20, 22, 14, 22, 22])
    for i, a in enumerate(A["app_wide"], start=2):
        _cell(ws4, i, 1, a["incident"])
        _cell(ws4, i, 2, a["opened"])
        _cell(ws4, i, 3, ", ".join(sorted(a["platforms"])))
        _cell(ws4, i, 4, a["gap"], "#,##0")
        _cell(ws4, i, 5, a["sessions"], "#,##0")
        _cell(ws4, i, 6, round(a["share"], 2), "0.00")
    ws4.freeze_panes = "A2"
    ws5 = wb.create_sheet("Notes")
    notes = ["Sources: experiment registry, install exports (service 1 for Q1, service 2 after), incident register, "
             "reconciliation summary and weekly review log, all as pulled on 12 October 2026.",
             "An experiment counts as one with a confirmed data incident when the register holds a confirmed entry "
             "against it (18 of 82); withdrawn suspicions count as experiments without one.",
             "Weekly review: an experiment counts as flagged if any review logged a flag against it, cleared or not; "
             "types are the causes the reconciliation established.",
             "App-wide incidents: the register repeats each incident's gap sessions on every platform row, so each "
             "incident's figure is taken once; its sessions are the sum of its rows.",
             "Hold line: 2.576 standard errors of the difference in gap-session share for 10,000 installs an arm, the "
             "per-install variance taken from both arms of the 64 experiments without a confirmed incident."]
    for i, t in enumerate(notes, start=1):
        c = ws5.cell(row=i, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name="Calibri", size=11)
    ws5.column_dimensions["A"].width = 110
    wb.properties.creator = "Mathilde Knudsen"
    wb.properties.lastModifiedBy = "Mathilde Knudsen"
    path = os.path.join(folder, XLSX)
    wb.save(path)
    postprocess.scrub(path, "Mathilde Knudsen", "2026-10-14")
    postprocess.finish_office(path, "2026-10-13T08:05:00Z", "2026-10-14T16:20:00Z", W.ORG)
    return path
