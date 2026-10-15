"""Independent verifier: reads only the files under target/ (csv, json, pyarrow, openpyxl and pypdf; no pandas, no
generator state) and recomputes the call, every check's record, the hold line, the weekly review by type and the
app-wide incidents, then compares them with figures.json."""
import csv
import json
import os
import re
from math import floor, sqrt

import pyarrow.parquet as pq
from openpyxl import load_workbook
from pypdf import PdfReader

Z = 2.5758293035489
STD_N = 10000
NAMES = {"gap": "Gap sessions", "dup": "Duplicate sessions", "late": "Late sessions", "epu": "Events per install",
         "share": "Treatment share"}
TYPES = ["telemetry loss", "duplicate delivery", "late delivery", "consent loss"]


def norm(s):
    return re.sub(r"\s+", " ", str(s)).strip().lower()


class Arm:
    __slots__ = ("n", "y", "yy", "e", "ee", "x", "xx", "xy")

    def __init__(self):
        self.n = 0
        self.y = self.yy = self.e = self.ee = 0.0
        self.x = {m: 0.0 for m in ("gap", "dup", "late")}
        self.xx = {m: 0.0 for m in ("gap", "dup", "late")}
        self.xy = {m: 0.0 for m in ("gap", "dup", "late")}

    def add(self, s, g, d, l, ev):
        self.n += 1
        self.y += s
        self.yy += s * s
        self.e += ev
        self.ee += ev * ev
        for m, x in (("gap", g), ("dup", d), ("late", l)):
            self.x[m] += x
            self.xx[m] += x * x
            self.xy[m] += x * s


def main(target, figures_path):
    fig = json.load(open(figures_path))
    p = lambda f: os.path.join(target, f)
    reg = json.load(open(p("experiment_registry_2026.json")))["experiments"]
    design = {e["key"]: e["design"]["treatment"] for e in reg}
    by_title = {norm(e["title"]): e["key"] for e in reg}
    wb = load_workbook(p("data_incident_register_2026.xlsx"), read_only=True)
    ws = wb["Register"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    regrows = [dict(zip(hdr, r)) for r in rows[1:]]
    inc = {by_title[norm(r["experiment"])] for r in regrows if r["scope"] == "experiment" and r["status"] == "confirmed"}
    arms = {}
    with open(p("install_summary_2026q1.csv"), newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            a = arms.setdefault((r["exp_id"], "control" if r["variant"] == "A" else "treatment"), Arm())
            a.add(int(r["n_sessions"]), int(r["sessions_seq_gap"]), int(r["sessions_dup_ids"]), int(r["sessions_late"]),
                  int(r["events"]))
    for q in ("install_summary_2026q2.parquet", "install_summary_2026q3.parquet"):
        t = pq.read_table(p(q)).to_pydict()
        for k, arm, s, g, d, l, ev in zip(t["experiment_key"], t["arm"], t["sessions"], t["gap_sessions"],
                                          t["duplicate_sessions"], t["late_sessions"], t["events_received"]):
            arms.setdefault((k, arm), Arm()).add(s, g, d, l, ev)
    keys = sorted({k for k, _ in arms})
    assert len(keys) == 82 and set(keys) == set(design)

    def ratio(a, m):
        n = a.n
        mx, my = a.x[m] / n, a.y / n
        vx = (a.xx[m] - n * mx * mx) / (n - 1)
        vy = (a.yy - n * my * my) / (n - 1)
        cxy = (a.xy[m] - n * mx * my) / (n - 1)
        r = mx / my
        return r, (vx - 2 * r * cxy + r * r * vy) / (my * my * n)

    held = {m: set() for m in NAMES}
    for k in keys:
        c, t = arms[(k, "control")], arms[(k, "treatment")]
        for m in ("gap", "dup", "late"):
            rc, vc = ratio(c, m)
            rt, vt = ratio(t, m)
            if (rt - rc) / sqrt(vc + vt) > Z:
                held[m].add(k)
        ec, et = c.e / c.n, t.e / t.n
        vec = (c.ee - c.n * ec * ec) / (c.n - 1)
        vet = (t.ee - t.n * et * et) / (t.n - 1)
        if (ec - et) / sqrt(vec / c.n + vet / t.n) > Z:
            held["epu"].add(k)
        n, pp = c.n + t.n, design[k]
        if (n * pp - t.n) / sqrt(n * pp * (1 - pp)) > Z:
            held["share"].add(k)
    none = [k for k in keys if k not in inc]
    cap = floor(0.02 * len(none) + 1e-9)
    rec = {m: (len(held[m] & inc), len(held[m] - inc)) for m in NAMES}
    okk = sorted([m for m in NAMES if rec[m][1] <= cap], key=lambda m: (-rec[m][0], rec[m][1]))
    pick = okk[0]
    # the line: per-install variance of the gap-session share over both arms of the experiments without an incident
    X = Y = XX = YY = XY = 0.0
    N = 0
    for k in none:
        for arm in ("control", "treatment"):
            a = arms[(k, arm)]
            X += a.x["gap"]; Y += a.y; XX += a.xx["gap"]; YY += a.yy; XY += a.xy["gap"]; N += a.n
    r = X / Y
    ss = XX - 2 * r * XY + r * r * YY
    s2 = ss / (N - 1) / (Y / N) ** 2
    line = 100 * Z * sqrt(2 * s2 / STD_N)
    # the weekly review by the cause the reconciliation established
    text = " ".join(pg.extract_text() for pg in PdfReader(p("reconciliation_summary_2026.pdf")).pages)
    cause = {}
    for m in re.finditer(r"(AV-26-\d{3})\b(.*?)(confirmed|not confirmed)\s*(telemetry loss|duplicate delivery|late delivery|consent loss)?",
                         text, re.S):
        if m.group(3) == "confirmed":
            cause[m.group(1)] = m.group(4)
    assert set(cause) == inc, (sorted(set(cause) ^ inc))
    flagged = set()
    with open(p("weekly_review_log_2026.csv"), newline="", encoding="utf-8") as f:
        for r_ in csv.DictReader(f):
            if r_["action"] == "flag":
                flagged.add(r_["experiment_key"])
    review = {t: [0, 0] for t in TYPES}
    for k in inc:
        review[cause[k]][0 if k in flagged else 1] += 1
    # the app-wide incidents
    aw = {}
    for r_ in regrows:
        if r_["scope"] == "app-wide":
            a = aw.setdefault(r_["incident_id"], {"gap": set(), "sessions": 0})
            a["gap"].add(int(r_["gap_sessions"]))
            a["sessions"] += int(r_["sessions_in_window"])
    # compare
    assert NAMES[pick] == fig["call"], (pick, fig["call"])
    assert "%.2f" % line == fig["line"], (line, fig["line"])
    for m in NAMES:
        assert (fig["record"][NAMES[m]]["incidents"], fig["record"][NAMES[m]]["none"]) == rec[m], (m, rec[m])
    assert NAMES[okk[1]] == fig["runner_up"]
    for t in TYPES:
        assert [fig["review"][t]["flagged"], fig["review"][t]["never"]] == review[t], (t, review[t])
    for iid, a in aw.items():
        assert len(a["gap"]) == 1
        g = a["gap"].pop()
        assert fig["app_wide"][iid]["gap"] == "{:,}".format(g)
        assert fig["app_wide"][iid]["share"] == "%.2f" % (100.0 * g / a["sessions"]), iid
    print("verify: call %s at %.4f points; records %s; review %s; %d app-wide incidents" % (pick, line, rec, review, len(aw)))


if __name__ == "__main__":
    import sys
    here = os.path.dirname(os.path.abspath(__file__))
    main(os.path.join(os.path.dirname(here), "target"), os.path.join(here, "figures.json"))
