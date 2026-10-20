"""The golden deliverables: the district leak methane inventory (xlsx) and the designation memo to the Council (PDF
with a chart of each district's methane per mile against the line), and generator/figures.json."""
import json
import os

from reportlab import rl_config
rl_config.invariant = 1
from reportlab.graphics.shapes import Drawing, Line, Rect, String                       # noqa: E402
from reportlab.lib import colors                                                        # noqa: E402
from reportlab.lib.pagesizes import letter                                              # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                         # noqa: E402
from reportlab.lib.units import inch                                                    # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402
from openpyxl import Workbook                                                           # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill                                # noqa: E402

import docs as DC   # noqa: E402
import methods as M  # noqa: E402
import pack as P    # noqa: E402
import world as W   # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
ORANGE, BLUE, INK = colors.HexColor("#c8553d"), colors.HexColor("#4f7ca8"), colors.HexColor("#222222")
WORDS = {0: "none", 1: "one", 2: "two", 3: "three", 4: "four", 5: "five", 6: "six", 7: "seven", 8: "eight"}
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
          "November", "December"]


def figures(x):
    f = x["fig"]
    rows = []
    for r in f["rows"]:
        rows.append(dict(district=r["district"], month=MONTHS[W.DISTRICTS[r["district"]][0] - 1], leaks=r["leaks"],
                         tonnes="%.1f" % r["tonnes"], miles="%.1f" % r["miles"], per_mile="%.2f" % r["per_mile"],
                         designated=r["designated"]))
    des = [r["district"] for r in rows if r["designated"]]
    below = sorted([r for r in f["rows"] if not r["designated"]], key=lambda r: -r["per_mile"])
    above = sorted([r for r in f["rows"] if r["designated"]], key=lambda r: r["per_mile"])
    g = dict(rows=rows, call=sorted(des), n_call=len(des), line="%.2f" % W.LINE,
             nearest_below=below[0]["district"], nearest_below_rate="%.2f" % below[0]["per_mile"],
             lowest_above=above[0]["district"], lowest_above_rate="%.2f" % above[0]["per_mile"],
             profile=["%.1f" % (100 * p) for p in f["profile"]], frost_share="%.1f" % (100 * sum(
                 f["profile"][i] for i in (11, 0, 1, 2))), sections=DC.SECTIONS,
             ask_a={d: int(round(v)) for d, v in f["ask_a"].items()}, ask_b=dict(f["ask_b"]),
             rates={k: "%.2f" % v for k, v in W.RATE.items()})
    g["districts"] = {r["district"]: dict(leaks=r["leaks"], tonnes=float(r["tonnes"]), per_mile=float(r["per_mile"]),
                                          designated=r["designated"]) for r in rows}
    return g


def write_workbook(path, g, x):
    wb = Workbook()
    ws = wb.active
    ws.title = "Districts"
    head = ["District", "Survey month", "Leaks found 2026", "Leak methane 2026 (t)", "CI/BS main (miles)",
            "Methane per mile (t)", "Designated", "2026 leak repair cost ($)", "Customer interruptions 2026"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4F7CA8")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in g["rows"]:
        ws.append([r["district"], r["month"], r["leaks"], float(r["tonnes"]), float(r["miles"]), float(r["per_mile"]),
                   "Yes" if r["designated"] else "No", g["ask_a"][r["district"]], g["ask_b"][r["district"]]])
    ws.append(["Total", None, sum(r["leaks"] for r in g["rows"]), round(sum(float(r["tonnes"]) for r in g["rows"]), 1),
               round(sum(float(r["miles"]) for r in g["rows"]), 1), None, "%d of 12" % g["n_call"],
               sum(g["ask_a"].values()), sum(g["ask_b"].values())])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        row[3].number_format = "0.0"
        row[4].number_format = "0.0"
        row[5].number_format = "0.00"
        row[7].number_format = "$#,##0"
        if row[6].value == "Yes":
            for c in row[:7]:
                c.font = Font(bold=True)
    for col, w in zip("ABCDEFGHI", [14, 12, 11, 12, 11, 11, 11, 14, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B2"
    ws2 = wb.create_sheet("Leak starts")
    ws2.append(["Month", "Share of odor-reported leak starts, 2016-2026 (%)"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for m, p in zip(MONTHS, g["profile"]):
        ws2.append([m, float(p)])
    ws2.column_dimensions["A"].width = 12
    ws2.column_dimensions["B"].width = 44
    ws3 = wb.create_sheet("Notes")
    for line in [
        "Leak methane (guidelines 2.2): each 2026 leak's Table 1 rate for its grade when found, times the hours from its "
        "start to its repair or to 31 December 2026.",
        "A leak found after an odor call started at the call (local time). A leak found by the survey started at an "
        "unrecorded time after its segment's 2025 survey; its start is the expected start over that interval with each "
        "day weighted by its month's share of the odor-reported leak starts (sheet Leak starts).",
        "Miles: in-service cast-iron and bare-steel segments in the GIS on 31 December 2026 (guidelines 2.3).",
        "Repair cost: each work order's cost of each type at the end of 2026 less at the end of 2025 (the cost lines "
        "carry the cost to date), by the leak's own district.",
        "Customer interruptions: each service address shut off for a leak repair, once per leak, its meters counted.",
        "Prepared by %s, %s." % (W.PEOPLE["engineer"][0], W.MEMO_DATE)]:
        ws3.append([line])
    ws3.column_dimensions["A"].width = 120
    wb.save(path)


def chart(g):
    rows = sorted(g["rows"], key=lambda r: -float(r["per_mile"]))
    w, h = 6.6 * inch, 3.3 * inch
    d = Drawing(w, h)
    left, bottom, top = 0.55 * inch, 0.75 * inch, 0.3 * inch
    hi = 4.5
    sy = lambda v: bottom + v / hi * (h - bottom - top)
    bw = (w - left - 10) / len(rows)
    for v in (0, 1, 2, 3, 4):
        d.add(Line(left, sy(v), w - 6, sy(v), strokeColor=colors.HexColor("#e2e5e8"), strokeWidth=0.5))
        d.add(String(left - 6, sy(v) - 3, "%.0f" % v if v else "0", fontName="Helvetica", fontSize=7,
                     textAnchor="end", fillColor=INK))
    for i, r in enumerate(rows):
        v = float(r["per_mile"])
        x0 = left + i * bw + 3
        col = ORANGE if r["designated"] else BLUE
        d.add(Rect(x0, bottom, bw - 6, sy(v) - bottom, fillColor=col, strokeColor=None))
        d.add(String(x0 + (bw - 6) / 2, sy(v) + 3, "%.2f" % v, fontName="Helvetica", fontSize=6.5,
                     textAnchor="middle", fillColor=INK))
        d.add(String(x0 + (bw - 6) / 2, bottom - 10, r["district"], fontName="Helvetica", fontSize=6.5,
                     textAnchor="middle", fillColor=INK))
        d.add(String(x0 + (bw - 6) / 2, bottom - 19, r["month"][:3], fontName="Helvetica", fontSize=6,
                     textAnchor="middle", fillColor=colors.HexColor("#666666")))
    d.add(Line(left, sy(W.LINE), w - 6, sy(W.LINE), strokeColor=INK, strokeWidth=1, strokeDashArray=[3, 2]))
    d.add(String(w - 8, sy(W.LINE) + 3, "Program line, %.2f t per mile" % W.LINE, fontName="Helvetica-Bold",
                 fontSize=7, textAnchor="end", fillColor=INK))
    d.add(String(left, h - 12, "Leak methane per mile of cast-iron and bare-steel main, 2026 (tonnes); district "
                 "survey month below each bar", fontName="Helvetica", fontSize=7.5, fillColor=INK))
    return d


def write_memo(path, g):
    body = ParagraphStyle("b", fontName="Times-Roman", fontSize=10.5, leading=13.6, spaceAfter=7)
    small = ParagraphStyle("s", fontName="Helvetica", fontSize=7.5, leading=9.5, textColor=colors.HexColor("#333333"))
    hd = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=BLUE)
    ct = ParagraphStyle("ct", fontName="Helvetica-Bold", fontSize=11, leading=14, spaceAfter=4)
    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.95 * inch, rightMargin=0.95 * inch,
                            topMargin=0.8 * inch, bottomMargin=0.8 * inch, title="Accelerated main replacement "
                            "designations 2027", author=W.PEOPLE["engineer"][0], creator=W.UTILITY,
                            producer=W.UTILITY, subject="Memo")

    def frame(c, d_):
        c.saveState()
        c.setFont("Helvetica", 7.5)
        c.drawString(0.95 * inch, 0.5 * inch, "%s, %s" % (W.UTILITY, W.DEPT))
        c.drawRightString(letter[0] - 0.95 * inch, 0.5 * inch, "Page %d" % d_.page)
        c.restoreState()
    rec = g["call"]
    s = [Paragraph("%s, %s" % (W.UTILITY, W.DEPT), hd), Spacer(1, 6),
         Table([["To:", "City Council, Budget Committee"], ["From:", "%s, %s" % W.PEOPLE["engineer"]],
                ["Date:", W.MEMO_DATE], ["Re:", "Accelerated Main Replacement Program 2027-2029: %s of 12 districts "
                                                "designated" % WORDS[g["n_call"]]]],
               colWidths=[0.7 * inch, 5.9 * inch],
               style=TableStyle([("FONT", (0, 0), (0, -1), "Helvetica-Bold", 9.5),
                                 ("FONT", (1, 0), (1, -1), "Helvetica", 9.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                                 ("LINEBELOW", (0, -1), (-1, -1), 0.6, colors.black)])),
         Spacer(1, 10),
         Paragraph("Under section %s of the program guidelines, %s districts are designated for 2027 to 2029: <b>%s</b>. "
                   "Each emitted %s tonnes of leak methane or more per mile of cast-iron and bare-steel main in 2026; the "
                   "lowest of them is %s at %s, and the nearest district below the line is %s at %s." % (
                       g["sections"]["designation"], WORDS[g["n_call"]], ", ".join(rec[:-1]) + " and " + rec[-1],
                       g["line"], g["lowest_above"], g["lowest_above_rate"], g["nearest_below"],
                       g["nearest_below_rate"]), body),
         Paragraph("Most of a surveyed leak's methane is emitted before the crew finds it, so the inventory has to "
                   "date each leak's start. Leaks reported by the public are found within days of starting, and %s per "
                   "cent of them started in December to March, when the ground freezes on the old mains. A leak found "
                   "on the yearly walk began somewhere in the year since the segment was last walked, and we dated it "
                   "by that seasonal pattern rather than at the midpoint of the year: a district walked in October or "
                   "November finds leaks that have been emitting since the winter, one walked in March or April finds "
                   "leaks a few weeks old. That is why Brennan Hill, Fenwick and Harlow, walked late in the year, are "
                   "on the list and spring-walked Ashby is not." % g["frost_share"], body),
         chart(g),
         Paragraph("Figure 1. Each district's 2026 leak methane per mile against the program line (guidelines %s). "
                   "Source: leak survey finds, odor call log and repair orders, 2026; GIS main segments at 31 December "
                   "2026." % g["sections"]["designation"], small), Spacer(1, 8)]
    t = [["District", "Leaks\nfound", "Methane\n(t)", "CI/BS\nmiles", "Per mile\n(t)", "Designated",
          "Leak repair\ncost 2026", "Customer\ninterruptions"]]
    for r in g["rows"]:
        t.append([r["district"], str(r["leaks"]), r["tonnes"], r["miles"], r["per_mile"],
                  "Yes" if r["designated"] else "No", "${:,}".format(g["ask_a"][r["district"]]),
                  str(g["ask_b"][r["district"]])])
    tab = Table(t, colWidths=[1.05 * inch, 0.55 * inch, 0.65 * inch, 0.6 * inch, 0.65 * inch, 0.75 * inch,
                              0.95 * inch, 0.85 * inch], repeatRows=1)
    tab.setStyle(TableStyle([("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 7.5), ("FONT", (0, 1), (-1, -1),
                                                                                 "Helvetica", 8),
                             ("ALIGN", (1, 0), (-1, -1), "RIGHT"), ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.black),
                             ("VALIGN", (0, 0), (-1, 0), "BOTTOM")]))
    s += [Paragraph("Table 1. Every district's inventory and the two figures section %s asks for the hearing." %
                    g["sections"]["reporting"], small), Spacer(1, 3), tab]
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


def build(gdir, x):
    import postprocess as PP
    if not os.path.exists(gdir):
        os.makedirs(gdir)
    g = figures(x)
    write_workbook(os.path.join(gdir, W.GOLDEN["workbook"]), g, x)
    PP.normalize_ooxml(os.path.join(gdir, W.GOLDEN["workbook"]), W.PEOPLE["engineer"][0], W.PEOPLE["engineer"][0],
                       "2027-04-02T14:10:00Z", "2027-04-17T20:51:44Z", excel=True, company=W.UTILITY)
    write_memo(os.path.join(gdir, W.GOLDEN["memo"]), g)
    PP.normalize_pdf(os.path.join(gdir, W.GOLDEN["memo"]), b"D:20270419163108-05'00'", b"D:20270419163108-05'00'")
    json.dump(g, open(os.path.join(HERE, "figures.json"), "w"), indent=1, sort_keys=True)
    return g
