"""The independent verifier: reads only target/ and golden/, recomputes the call, every district's figures and both
asks, and checks the golden workbook and memo carry them."""
import os
import re
import sys

from openpyxl import load_workbook
from pypdf import PdfReader

import bundle as B

HERE = os.path.dirname(os.path.abspath(__file__))


def main(target, golden):
    fails = []

    def ok(label, cond):
        print("%s  verify: %s" % ("ok  " if cond else "FAIL", label))
        if not cond:
            fails.append(label)
    g = re.sub(r"\s+", " ", " ".join(p.extract_text() for p in PdfReader(B.find(target, "amr_program")).pages))
    ok("pdf read back: guidelines 2.1 to 2.3 and Table 1", "3.30 tonnes or more" in g and
       "began when the call was received" in g and "1.62" in g and "in-service cast-iron and bare-steel" in g)
    inv = B.inventory(target)
    call = B.designated(inv)
    ok("the call: %s" % ", ".join(call), call == ["Brennan Hill", "Elmwood", "Fenwick", "Harlow", "Ivers"])
    wb = load_workbook(B.find(golden, "leak_methane_inventory"), read_only=True)
    rows = {r[0]: r for i, r in enumerate(wb["Districts"].iter_rows(values_only=True)) if i and r[0] != "Total"}
    a, b = B.ask_a(target), B.ask_b(target)
    for d, (nl, t, mi, pm) in sorted(inv.items()):
        w = rows[d]
        ok("%s: %d leaks, %.1f t, %.2f per mile, $%d, %d interruptions" % (d, nl, t, pm, round(a[d]), b[d]),
           w[2] == nl and "%.1f" % w[3] == "%.1f" % t and "%.2f" % w[5] == "%.2f" % pm and
           (w[6] == "Yes") == (d in call) and w[7] == round(a[d]) and w[8] == b[d])
    m = re.sub(r"\s+", " ", " ".join(p.extract_text() for p in PdfReader(B.find(golden, "main_replacement")).pages))
    ok("the memo designates the five", "Brennan Hill, Elmwood, Fenwick, Harlow and Ivers" in m)
    print("verify: %d failures" % len(fails))
    return not fails


if __name__ == "__main__":
    task = os.path.dirname(HERE)
    sys.exit(0 if main(os.path.join(task, "target"), os.path.join(task, "golden")) else 1)
