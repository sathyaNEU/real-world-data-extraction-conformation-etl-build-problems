"""Independent verifier for task115. Reads target/, golden/, the two archives, prompt.md and submission.md;
imports nothing from generator/ and shares no code with it.

Recomputes every graded figure on its own path (plain csv and dicts, no pandas), then checks the goldens, the
write-up, the archives and the house rules against it.
    python3 verify_pack.py
"""
import csv
import datetime as dt
import hashlib
import io
import logging
import math
import re
import sys
import zipfile
from collections import defaultdict
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent
T = ROOT / "target"
GD = ROOT / "golden"
BAD = []
FEED = "TullanInsight_tea_store_week_2024W41-2026W40.csv"
DEL160 = "211402"
MENU = [("Save a third", "pct", 1 / 3), ("Save 40%", "pct", 0.40), ("Half price", "pct", 0.50),
        ("€2.40 off", "abs", 2.40)]
REGIONS = ["Dublin & North East", "South", "West & North West", "Midlands & South East"]


def say(ok, what, info=""):
    print(("PASS " if ok else "FAIL ") + what + ("  " + str(info) if info != "" else ""))
    if not ok:
        BAD.append(what)


def d(s):
    return dt.date.fromisoformat(s)


def cents_floor(x):
    return math.floor(round(x * 100, 6)) / 100


def half_up(x, step=1.0):
    return math.floor(x / step + 0.5) * step


# ---------------------------------------------------------------------------- the files
feed = {}
by_item_week = defaultdict(list)
with open(T / FEED, encoding="utf-8") as fh:
    for r in csv.DictReader(fh):
        key = (r["store_no"], d(r["week_commencing"]), r["item_no"])
        rec = (int(r["units"]), float(r["regular_price_eur"]),
               float(r["deal_price_eur"]) if r["deal_price_eur"] else None, r["deal_ref"])
        feed[key] = rec
        by_item_week[(r["item_no"], key[1])].append((key[0],) + rec)
NROWS = len(feed)
weeks = sorted({k[1] for k in feed})
items = {}
with open(T / "tullan_tea_range_2026.csv", encoding="utf-8") as fh:
    for r in csv.DictReader(fh):
        items[r["item_no"]] = r
own_by_bags = {r["bag_count"]: i for i, r in items.items() if r["supplier"] == "Tullan Grocers"}
OL160 = own_by_bags["160"]
events = []
with open(T / "big_deals_calendar_2024-2026.csv", encoding="utf-8") as fh:
    for r in csv.DictReader(fh):
        if r["category"] == "Tea":
            v = float(r["deal_value"])
            kind = "pct" if r["deal_type"] == "PCT" else "abs"
            v = (1 / 3 if abs(v - 33.33) < 0.01 else v / 100) if kind == "pct" else v
            events.append((r["cycle"], d(r["start_week"]), r["item_no"], kind, v))
brief = Document(str(T / "Big_Deals_Spring_2027_supplier_brief.docx"))
brief_text = "\n".join(p.text for p in brief.paragraphs)
appb = set()
for t in brief.tables:
    rows = [[c.text.strip() for c in r.cells] for r in t.rows]
    if rows[0][0] == "Store no":
        appb = {r[0] for r in rows[1:]}
m = re.search(r"weeks commencing (\d+) (\w+) to (\d+) (\w+) (\d{4})", brief_text)
b_from = dt.datetime.strptime("%s %s %s" % (m.group(1), m.group(2), m.group(5)), "%d %B %Y").date()
b_to = dt.datetime.strptime("%s %s %s" % (m.group(3), m.group(4), m.group(5)), "%d %B %Y").date()
FLOOR = 2.5 if "two and a half times" in brief_text else None
say(len(appb) == 11 and FLOOR == 2.5 and (b_to - b_from).days == 84, "brief read: Appendix B, floor, base weeks",
    (len(appb), FLOOR, b_from, b_to))

# ---------------------------------------------------------------------------- regions in March 2027
wb = load_workbook(T / "tullan_members_jan2026.xlsx", read_only=True)
mrows = list(wb["Members Jan 2026"].iter_rows(values_only=True))
hdr = list(mrows[0])
region = {str(r[0]).zfill(4): r[hdr.index("Trading region")] for r in mrows[1:] if r[0]}
names = {str(r[0]).zfill(4): r[1] for r in mrows[1:] if r[0]}
scanning = {str(r[0]).zfill(4): r[hdr.index("Scan data")] for r in mrows[1:] if r[0]}
area = Document(str(T / "tullan_area_structure_from_6_July_2026.docx"))
for t in area.tables:
    for r in t.rows[1:]:
        c = [x.text.strip() for x in r.cells]
        if dt.datetime.strptime(c[4], "%d %B %Y").date() <= dt.date(2027, 3, 1):
            region[c[0]] = c[3]
news = Document(str(T / "member_news_2026.docx"))
for t in news.tables:
    for r in t.rows[1:]:
        c = [x.text.strip() for x in r.cells]
        if c[1] == "New store number":
            new = re.search(r"store (\d{4})", c[2]).group(1)
            region[new] = region[c[0]]
            names[new] = names[c[0]]
        elif c[1] == "Opened":
            region[c[0]] = re.search(r"trading region (.+?)\.$", c[2]).group(1)
            names[c[0]] = re.search(r"^(Tullan [A-Za-z ]+?) opened", c[2]).group(1)

# ---------------------------------------------------------------------------- deal lift and the switch


def pre_base(item, start, n=13):
    pre = {start - dt.timedelta(days=7 * k) for k in range(1, n + 1)}
    acc = defaultdict(list)
    for w in pre:
        for s, u, reg, dp, ref in by_item_week.get((item, w), []):
            if dp is None:
                acc[s].append(u)
    return {s: sum(v) / len(v) for s, v in acc.items()}


num, den = defaultdict(float), defaultdict(float)
ol_lost = ol_norm = 0.0
cross_cycles = set()
crossing = []          # (depth, deal-line units, deal-line base, own-label base) per crossing store-week
for code, start, item, kind, v in events:
    pair = own_by_bags.get(items[item]["bag_count"])
    base = pre_base(item, start)
    obase = pre_base(pair, start) if pair else {}
    for k in range(3):
        w = start + dt.timedelta(days=7 * k)
        for s, u, reg, dp, ref in by_item_week.get((item, w), []):
            if dp is None or ref != code or s not in base:
                continue
            crosses = pair is not None and (s, w, pair) in feed and dp <= feed[(s, w, pair)][1] + 1e-9
            if crosses:
                cross_cycles.add(code)
                if s in obase:
                    ol_lost += feed[(s, w, pair)][0]
                    ol_norm += obase[s]
                    crossing.append((v if kind == "pct" else v / reg, u, base[s], obase[s]))
            elif kind == "pct":
                num[v] += u
                den[v] += base[s]
lift = {v: num[v] / den[v] - 1 for v in num}
loss = 1 - ol_lost / ol_norm
pts = sorted([(0.0, 0.0)] + list(lift.items()))


def lift_at(x):
    for (x0, y0), (x1, y1) in zip(pts, pts[1:]):
        if x <= x1 + 1e-12:
            return y0 + (y1 - y0) * (x - x0) / (x1 - x0)
    (x0, y0), (x1, y1) = pts[-2], pts[-1]
    return y1 + (y1 - y0) * (x - x1) / (x1 - x0)


# the switch as the deal line shows it: its units over base x (1 + lift at its depth), per own-label base
share = sum(u - b * (1 + lift_at(dpt)) for dpt, u, b, o in crossing) / sum(o for *_, o in crossing)
say(len(cross_cycles) == 6 and abs(share - 0.70) < 0.015 and abs(loss - 0.35) < 0.01,
    "switch measured on six crossing cycles: deal line gains twice what own label loses", (round(share, 4),
                                                                                         round(loss, 4)))

# ---------------------------------------------------------------------------- the March forecast
bw = [w for w in weeks if b_from <= w <= b_to]
last = max(weeks)
units_b, units_o = defaultdict(list), defaultdict(list)
for w in bw:
    for s, u, reg, dp, ref in by_item_week.get((DEL160, w), []):
        if dp is None:
            units_b[s].append(u)
    for s, u, reg, dp, ref in by_item_week.get((OL160, w), []):
        if dp is None:
            units_o[s].append(u)
stores = sorted(s for s in units_b if s not in appb)
fc = {}
for s in stores:
    reg, olp = feed[(s, last, DEL160)][1], feed[(s, last, OL160)][1]
    b, o = sum(units_b[s]) / len(units_b[s]), sum(units_o[s]) / len(units_o[s])
    row = {"base": b, "reg": reg}
    for name, kind, v in MENU:
        dp = cents_floor(reg * (1 - v)) if kind == "pct" else cents_floor(reg - v)
        depth = v if kind == "pct" else v / reg
        row[name] = (b * (1 + lift_at(depth)) + (share * o if dp <= olp + 1e-9 else 0.0), dp)
    fc[s] = row
B = sum(r["base"] for r in fc.values())
chain = {n: sum(r[n][0] for r in fc.values()) / B for n, _, _ in MENU}
reg_base = {g: sum(fc[s]["base"] for s in fc if region[s] == g) for g in REGIONS}
reg_mult = {g: {n: sum(fc[s][n][0] for s in fc if region[s] == g) / reg_base[g] for n, _, _ in MENU} for g in REGIONS}
packs = {g: 3 * sum(fc[s]["Save 40%"][0] for s in fc if region[s] == g) for g in REGIONS}
packs_all = 3 * sum(r["Save 40%"][0] for r in fc.values())
fund = {n: 3 * sum((r["reg"] - r[n][1]) * r[n][0] for r in fc.values()) for n, _, _ in MENU}
clear = [n for n, _, _ in MENU if chain[n] >= FLOOR]
pick = min(clear, key=lambda n: fund[n]) if clear else "sit out"
say(pick == "Save 40%", "cheapest deal clearing the floor", (pick, {n: round(x, 4) for n, x in chain.items()}))
say(len(stores) == 169 and all(region.get(s) in REGIONS for s in stores), "169 participating stores, all placed",
    len(stores))

# ---------------------------------------------------------------------------- the 2026 billing
dn = load_workbook(T / "tullan_debit_notes_Delahunt_2026.xlsx", read_only=True)
drows = list(dn.active.iter_rows(values_only=True))
dh = list(drows[0])
invoiced = defaultdict(float)
for r in drows[1:]:
    if r[0]:
        invoiced[r[dh.index("Cycle")]] += float(r[dh.index("Net (EUR)")])
supported = defaultdict(float)
for (s, w, i), (u, reg, dp, ref) in feed.items():
    if i == DEL160 and ref in invoiced and dp is not None:
        supported[ref] += round((reg - dp) * u, 2)
with open(T / "tullan_manual_returns_2026.csv", encoding="utf-8") as fh:
    for r in csv.DictReader(fh):
        if r["item_no"] == DEL160 and r["deal_ref"] in invoiced:
            supported[r["deal_ref"]] += (float(r["regular_price_eur"]) - float(r["deal_price_eur"])) * int(
                r["deal_units"])
claims = {c: (half_up(invoiced[c]), half_up(supported[c]), half_up(invoiced[c] - supported[c])) for c in invoiced}

# ---------------------------------------------------------------------------- expected answers
one = lambda x: "%.1f" % x
EXP = {
    "chain": {n: one(chain[n]) for n, _, _ in MENU},
    "region": {g: {n: one(reg_mult[g][n]) for n, _, _ in MENU} for g in REGIONS},
    "base": {g: int(half_up(reg_base[g])) for g in REGIONS},
    "base_total": int(half_up(B)),
    "packs": {g: int(half_up(packs[g], 1000)) for g in REGIONS},
    "packs_total": int(half_up(packs_all, 1000)),
    "claims": claims,
}
print("expected", EXP)
say(EXP["chain"] == {"Save a third": "1.7", "Save 40%": "2.7", "Half price": "4.4", "€2.40 off": "1.7"},
    "chain multiples 1.7, 2.7, 4.4, 1.7")
say(sum(EXP["base"].values()) == EXP["base_total"] and sum(EXP["packs"].values()) == EXP["packs_total"],
    "regional bases and packs add to their totals at the stated rounding")


def edge(x, step):
    k = x / step
    return abs(k - math.floor(k) - 0.5)


margins = [edge(chain[n], 0.1) for n, _, _ in MENU] + [edge(reg_mult[g][n], 0.1) for g in REGIONS for n, _, _ in MENU]
say(min(margins) > 0.04, "every multiple at least 0.004 inside its bin", round(min(margins) * 0.1, 4))
say(min([edge(B, 1)] + [edge(reg_base[g], 1) for g in REGIONS]) > 0.15, "bases clear of their whole-pack edges")
say(min([edge(packs_all, 1000)] + [edge(packs[g], 1000) for g in REGIONS]) > 0.05, "packs clear of their bins")

# ---------------------------------------------------------------------------- golden: the workbook
wb = load_workbook(GD / "march_offer_workings.xlsx")
of = wb["Offer form"]
form = {of.cell(row=r, column=1).value: (of.cell(row=r, column=2).value, of.cell(row=r, column=3).value)
        for r in range(1, of.max_row + 1)}
say(form["Mechanic"][0] == "Save 40%" and "%.1f" % form["Forecast multiple of base"][0] == EXP["chain"]["Save 40%"],
    "offer form: mechanic and multiple")
say(all(form[g] == (EXP["base"][g], EXP["packs"][g]) for g in REGIONS) and
    form["Total"] == (EXP["base_total"], EXP["packs_total"]), "offer form: base and packs by region and total")
gr = wb["Forecast by deal"]
grid_ok = True
for r in range(4, 8):
    n = gr.cell(row=r, column=1).value
    vals = ["%.1f" % gr.cell(row=r, column=c).value for c in range(2, 7)]
    grid_ok &= vals == [EXP["chain"][n]] + [EXP["region"][g][n] for g in REGIONS]
say(grid_ok, "workbook grid matches every recomputed multiple")
bc = wb["2026 billing check"]
bill = {bc.cell(row=r, column=1).value: tuple(bc.cell(row=r, column=c).value for c in (3, 4, 5)) for r in (4, 5)}
say(all(bill[c] == tuple(int(x) for x in claims[c]) for c in claims), "billing check matches", bill)
sw = wb["Store workings"]
tot = [sw.cell(row=sw.max_row, column=c).value for c in (6, 13)]
say(abs(tot[0] - B) < 0.05 and abs(3 * tot[1] - packs_all) < 0.5, "store workings total ties", tot)

# ---------------------------------------------------------------------------- golden: the paper
cap = io.StringIO()
h = logging.StreamHandler(cap)
logging.getLogger("pypdf").addHandler(h)
paper = " ".join(p.extract_text() for p in PdfReader(GD / "march_offer_paper.pdf").pages)
logging.getLogger("pypdf").removeHandler(h)
flat = re.sub(r"\s+", " ", paper)
say(not cap.getvalue().strip(), "paper parses without errors")
say(re.search(r"offer Tullan Save 40% on Original 160s.*forecast " + re.escape(EXP["chain"]["Save 40%"]) +
              r" times base", flat) is not None, "paper opens on the deal and its multiple")
rows_ok = True
for n, _, _ in MENU:
    want = " ".join([n, EXP["chain"][n]] + [EXP["region"][g][n] for g in REGIONS])
    rows_ok &= want in flat
say(rows_ok, "paper grid matches every recomputed multiple")
say("about {:,} packs".format(EXP["packs_total"]) in flat and "{:,} packs a week".format(EXP["base_total"]) in flat,
    "paper packs and base")
say(all(("€{:,}".format(int(claims[c][2]))) in flat for c in claims), "paper billing differences")
say(len(PdfReader(GD / "march_offer_paper.pdf").pages) == 1, "paper is one page")

# ---------------------------------------------------------------------------- submission.md
sub = (ROOT / "submission.md").read_text(encoding="utf-8")
blocks = re.split(r"\n## ", sub)
say([b.split("\n")[0] for b in blocks[1:]] == ["Tags", "1. Final Recommendation", "2. Critical Components",
                                               "3. Step-by-Step Solution", "4. Deliverable Answers"],
    "submission carries exactly the five blocks")
for n, _, _ in MENU:
    line = "%s: all stores %s, " % (n, EXP["chain"][n]) + ", ".join("%s %s" % (g, EXP["region"][g][n]) for g in REGIONS)
    say(line in sub, "submission grid row: %s" % n)
say("Dublin & North East {:,}, South {:,}, West & North West {:,}, Midlands & South East {:,}, total {:,}".format(
    *[EXP["base"][g] for g in REGIONS], EXP["base_total"]) in sub, "submission bases")
say("Dublin & North East {:,}, South {:,}, West & North West {:,}, Midlands & South East {:,}".format(
    *[EXP["packs"][g] for g in REGIONS]) in sub, "submission regional packs")
for c, (i, s, x) in claims.items():
    say("%s (Save %s): invoiced {:,}, supported {:,}, difference {:,}".format(int(i), int(s), int(x))
        % (c, "20%" if c == "BD26-04" else "25%") in sub, "submission billing %s" % c)
# every number in the write-up is owed: sweep by unit class against what this verifier computed
allowed = set()
for n in EXP["chain"].values():
    allowed.add(n)
for g in REGIONS:
    allowed |= set(EXP["region"][g].values())
allowed |= {"{:,}".format(v) for v in list(EXP["base"].values()) + [EXP["base_total"]]}
allowed |= {"{:,}".format(v) for v in list(EXP["packs"].values()) + [EXP["packs_total"]]}
allowed |= {"{:,}".format(int(x)) for c in claims.values() for x in c}
prices = set()
for r in fc.values():
    prices.add("%.2f" % r["reg"])
    for n, _, _ in MENU:
        prices.add("%.2f" % r[n][1])
for s in stores:
    prices.add("%.2f" % feed[(s, last, OL160)][1])
allowed |= prices | {"3.79", "2.40", "0.7", "2.5", "1.7"}
allowed |= {"115", "2027", "2026", "2024", "13", "29", "21", "22", "6", "5", "1", "9", "14", "0147", "0388", "0119",
            "26-04", "26-13", "160", "20", "25", "40", "211402", "2", "3", "4", "8", "10"}
swept = re.sub(r"^\s*\d+\. ", "", sub, flags=re.M)                       # list numbering
swept = re.sub(r"\b(BD\d\d-\d\d|DN-\d\d-\d{4}|[\w-]+\.(?:csv|xlsx|docx|pdf|txt))\b", "", swept)
nums = re.findall(r"(?<![\w.])(\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?)(?![\d])", swept)
stray = sorted({x for x in nums if x not in allowed})
say(not stray, "every number in submission.md is owed or a stated parameter", stray)
say("—" not in sub and "—" not in (ROOT / "prompt.md").read_text() and "—" not in flat,
    "no em dashes in the write-up, the prompt or the paper")
say(abs(share - 0.7) < 0.05 and "**0.7**" in sub, "the 0.7 of own-label base in the write-up recomputes",
    round(share, 4))
say(abs(1 + lift_at(0.40) - 1.7) < 0.05, "Save 40% alone lifts a line to 1.7 times base", round(1 + lift_at(0.40), 4))

# ---------------------------------------------------------------------------- the prompt
prompt = (ROOT / "prompt.md").read_text(encoding="utf-8")
for pin in ("one decimal place", "nearest thousand", "whole euro", "whole packs a week", "ex VAT", "2.5"):
    say(pin in prompt or (pin == "2.5" and "two and a half times" in prompt), "prompt pins %s" % pin)
named = set(re.findall(r"\b([a-z_]+\.(?:pdf|xlsx))\b", prompt))
say(named == {p.name for p in GD.iterdir() if p.is_file()}, "golden holds exactly the files the prompt names", named)

# ---------------------------------------------------------------------------- house checks
files = sorted(p for p in T.iterdir() if p.is_file())
say(len(files) >= 10 and len({p.suffix for p in files}) >= 3 and NROWS >= 25000, "input gates",
    (len(files), sorted({p.suffix for p in files}), NROWS))
readme = (T / "README_tullan_spring_2027.txt").read_text()
listed = set(re.findall(r"^\s+(\S+\.(?:csv|pdf|docx|xlsx|txt))", readme, re.M))
say(listed | {"README_tullan_spring_2027.txt"} == {p.name for p in files}, "H9 README lists every shipped file")
SIG = [b"python-docx", b"openpyxl", b"Openpyxl", b"reportlab", b"ReportLab", b"matplotlib", b"Matplotlib"]
meta_bad, date_bad = [], []
for p in files + sorted(GD.iterdir()):
    if p.suffix in (".docx", ".xlsx"):
        with zipfile.ZipFile(p) as z:
            blob = b"".join(z.read(n) for n in z.namelist() if n.startswith("docProps/"))
            stamps = [x.decode()[:10] for x in re.findall(rb"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", blob)]
    elif p.suffix == ".pdf":
        blob = p.read_bytes()
        stamps = ["%s-%s-%s" % (x[2:6].decode(), x[6:8].decode(), x[8:10].decode())
                  for x in re.findall(rb"D:\d{14}", blob)]
    else:
        continue
    if any(s in blob for s in SIG):
        meta_bad.append(p.name)
    if any(not ("2024-01-01" <= s <= "2026-10-31") for s in stamps) or not stamps:
        date_bad.append((p.name, stamps))
say(not meta_bad, "H1 no writer's name in any container", meta_bad)
say(not date_bad, "H1 every container date inside the setting", date_bad)
for arch, folder in (("target.zip", T), ("golden.zip", GD)):
    a = ROOT / arch
    if not a.exists():
        say(False, "H4 %s exists" % arch)
        continue
    with zipfile.ZipFile(a) as z:
        names = z.namelist()
        same = sorted(names) == sorted(p.name for p in folder.iterdir() if p.is_file()) and all(
            z.read(n) == (folder / n).read_bytes() for n in names)
        say(same and not any("/" in n or n.startswith(("__MACOSX", ".")) for n in names),
            "H4/H7 %s is flat and identical to %s/" % (arch, folder.name), "%d members, sha256 %s" % (
                len(names), hashlib.sha256(a.read_bytes()).hexdigest()[:16]))
tree = [p for p in ROOT.rglob("*") if p.is_file()]
say(sum(p.name == "submission.md" for p in tree) == 1 and sum(p.name == "prompt.md" for p in tree) == 1,
    "H11 one submission.md and one prompt.md under the task")
say(not [p for p in tree if re.search(r"(DESIGN_V\d|_v\d+_snapshot|snapshot\.zip|\.bak$|\.pre-scrub$|\.tmp$)",
                                      p.name)], "no backups, snapshots or versioned copies")
say(max(weeks) <= dt.date(2026, 9, 28) and "Prepared 6 October 2026" in " ".join(
    p.extract_text() for p in PdfReader(T / "big_deals_review_tea_2025-26.pdf").pages),
    "H16 the review is dated after the last cycle it reports")
texts = []
for p in files:
    if p.suffix in (".csv", ".txt"):
        texts.append(p.read_text(errors="ignore") if p.name != FEED else "")
    elif p.suffix == ".pdf":
        texts.append(" ".join(x.extract_text() for x in PdfReader(p).pages))
    elif p.suffix == ".docx":
        dd = Document(str(p))
        texts.append(" ".join(x.text for x in dd.paragraphs) + " ".join(c.text for t in dd.tables for r in t.rows
                                                                       for c in r.cells))
    elif p.suffix == ".xlsx":
        w2 = load_workbook(p, read_only=True)
        texts.append(" ".join(str(c) for ws in w2.worksheets for r in ws.iter_rows(values_only=True) for c in r if c))
alltext = " ".join(texts) + flat + sub + prompt
retired = [n for n in ("Abbeyfield", "ABBEYFIELD", "Oatfield", "OATFIELD", "Andrea Moloney") if n in alltext]
say(not retired, "H21 no retired name survives", retired)
for term in ("own brand", "Appendix B", "Save 40%", "€2.40 off", "Half price", "Save a third", "trading region"):
    say(term.lower() in " ".join(texts).lower(), "Gate B: '%s' occurs in an input" % term)

print("\n%d failures" % len(BAD))
for b in BAD:
    print("  FAIL", b)
sys.exit(1 if BAD else 0)
