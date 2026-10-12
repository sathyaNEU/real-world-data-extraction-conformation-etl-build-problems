"""The golden deliverables, built from target/ only: march_offer_paper.pdf and march_offer_workings.xlsx.

    python3 deliver.py            # after build.py
"""
import io
import math
from collections import Counter, defaultdict
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab import rl_config
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import golden as GD
import seal

rl_config.invariant = 1          # fixed dates and /ID; seal.py writes the in-fiction save time

ROOT = Path(__file__).resolve().parent.parent
TARGET = ROOT / "target"
OUT = ROOT / "golden"
REGIONS = ["Dublin & North East", "South", "West & North West", "Midlands & South East"]
DEALS = ["Save a third", "Save 40%", "Half price", "€2.40 off"]
OFFER = "Save 40%"
NAVY, SLATE, RED = "#1f4e79", "#aebccb", "#b03a2e"


def one(x):
    return "%.1f" % x


def thousand(x):
    return "{:,}".format(int(math.floor(x / 1000 + 0.5)) * 1000)


def whole(x):
    return "{:,}".format(int(math.floor(x + 0.5)))


def figures():
    g = GD.golden(TARGET)
    fc = g["fc"]
    fund = {d: 3 * sum((r["reg"] - r[d][1]) * r[d][0] for r in fc.values()) for d in DEALS}
    cleared = [d for d in DEALS if g["chain"][d] >= GD.FLOOR]
    assert min(cleared, key=lambda d: fund[d]) == OFFER
    assert 1.8 <= fund["Half price"] / fund[OFFER] < 2.2          # "about twice as much to fund"
    assert 0.30 <= g["loss_share"] <= 0.37 and 1.8 <= g["share"] / g["loss_share"] <= 2.2   # "about a third", "twice"
    # one own-brand price for each Delahunt price, and the deal prices at each
    prices = defaultdict(set)
    for r in fc.values():
        prices[r["reg"]].add(r["ol_reg"])
    assert all(len(v) == 1 for v in prices.values())
    ladder = [(reg, next(iter(prices[reg])), {d: GD.deal_price(reg, k, v) for d, k, v in GD.MENU})
              for reg in sorted(prices, reverse=True)]
    top = ladder[0][0]
    at_top = Counter(g["regions"][s] for s, r in fc.items() if r["reg"] == top)
    assert at_top["Dublin & North East"] > sum(at_top.values()) / 2     # "most of the EUR 6.99 stores"
    wb = load_workbook(TARGET / GD.MEMBERS_F, read_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    hdr = list(rows[0])
    names = {str(r[hdr.index("Store no")]).zfill(4): r[hdr.index("Store name")] for r in rows[1:]}
    names["0388"] = names.pop("0147")
    names["0119"] = "Tullan Swords Rivervalley"
    rv = load_workbook(TARGET / "delahunt_deal_reviews_2025-26.xlsx", read_only=True)
    rr = list(rv.active.iter_rows(values_only=True))
    own = [r[rr[0].index("Multiple of base")] for r in rr[1:] if r[0]]
    dn = load_workbook(TARGET / GD.DN_F, read_only=True)
    dr = list(dn.active.iter_rows(values_only=True))
    notes = [dict(zip(dr[0], r)) for r in dr[1:] if r[0]]
    return g, fund, ladder, names, own, notes


# ---------------------------------------------------------------- the chart
def chart(g):
    arial = [f for f in font_manager.findSystemFonts() if Path(f).name in ("Arial.ttf", "Arial Bold.ttf")]
    for f in arial:
        font_manager.fontManager.addfont(f)
    plt.rcParams.update({"font.family": "Arial" if arial else "DejaVu Sans", "font.size": 9})
    fig, ax = plt.subplots(figsize=(6.6, 2.29), dpi=200)
    vals = [g["chain"][d] for d in DEALS]
    bars = ax.bar(DEALS, vals, color=[NAVY if d == OFFER else SLATE for d in DEALS], width=0.56, zorder=2)
    ax.axhline(GD.FLOOR, color=RED, lw=1.2, ls=(0, (5, 3)), zorder=3)
    ax.text(3.45, GD.FLOOR + 0.08, "Tullan's floor: 2.5x base", color=RED, ha="right", va="bottom", fontsize=8)
    for b, v, d in zip(bars, vals, DEALS):
        ax.text(b.get_x() + b.get_width() / 2, v + 0.07, one(v) + "x", ha="center", va="bottom", fontsize=9,
                fontweight="bold" if d == OFFER else "normal", color=NAVY if d == OFFER else "#3a3a3a")
    ax.set_xticks(range(len(DEALS)))
    ax.set_xticklabels([d + "\n(our offer)" if d == OFFER else d for d in DEALS])
    for lab, d in zip(ax.get_xticklabels(), DEALS):
        if d == OFFER:
            lab.set_color(NAVY)
            lab.set_fontweight("bold")
    ax.set_ylim(0, 5.0)
    ax.set_yticks([0, 1, 2, 3, 4, 5])
    ax.set_yticklabels(["0", "1x", "2x", "3x", "4x", "5x"])
    ax.set_ylabel("Forecast multiple of base")
    ax.set_title("Save 40% is the cheapest deal that clears Tullan's floor", fontsize=10, loc="left",
                 fontweight="bold", color="#222222")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    for s in ("left", "bottom"):
        ax.spines[s].set_color("#8a8a8a")
    ax.tick_params(colors="#3a3a3a", length=0)
    ax.grid(axis="y", color="#e6e6e6", lw=0.6, zorder=0)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------- the paper
def footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(colors.HexColor("#666666"))
    canvas.drawString(17 * mm, 9 * mm, "Delahunt Tea Ltd, Revenue Growth. Commercial in confidence.")
    canvas.drawRightString(A4[0] - 17 * mm, 9 * mm, "Page %d" % doc.page)
    canvas.restoreState()


def paper(g, fund, ladder, own, path, png):
    body = ParagraphStyle("b", fontName="Helvetica", fontSize=9.1, leading=12.0, spaceAfter=4)
    lead = ParagraphStyle("l", parent=body, fontName="Helvetica-Bold", fontSize=10.2, leading=13.6, spaceAfter=5)
    h2 = ParagraphStyle("h2", parent=body, fontName="Helvetica-Bold", fontSize=9.8, spaceBefore=5, spaceAfter=2.5,
                        textColor=colors.HexColor(NAVY))
    small = ParagraphStyle("s", parent=body, fontSize=7.6, leading=9.6, textColor=colors.HexColor("#555555"),
                           spaceAfter=3)
    title = ParagraphStyle("t", parent=body, fontName="Helvetica-Bold", fontSize=14, leading=17.5, spaceAfter=4)
    cell = ParagraphStyle("c", fontName="Helvetica", fontSize=8.5, leading=10.2)
    cellb = ParagraphStyle("cb", parent=cell, fontName="Helvetica-Bold")
    cellr = ParagraphStyle("cr", parent=cell, alignment=2)
    cellrb = ParagraphStyle("crb", parent=cellr, fontName="Helvetica-Bold")
    (p_hi, o_hi, d_hi), (p_mid, o_mid, d_mid), (p_lo, o_lo, d_lo) = ladder
    ch = {d: one(g["chain"][d]) for d in DEALS}
    rest = [g["region"][r][OFFER] for r in REGIONS[1:]]
    st = [Paragraph("DELAHUNT TEA &nbsp;|&nbsp; Revenue Growth &nbsp;|&nbsp; For the leadership team, "
                    "22 October 2026", small), Spacer(1, 1.5 * mm),
          Paragraph("Save 40% gets our 160s into Tullan's March leaflet", title),
          Paragraph("Recommendation: offer Tullan Save 40%% on Original 160s for Big Deals BD27-03, 1 to 21 March "
                    "2027. We forecast "
                    "%s times base over the three weeks against Tullan's floor of 2.5, and Save 40%% is the "
                    "cheapest of the four deals on the menu that clears it." % ch[OFFER], lead),
          Paragraph("That is about %s packs over the three weeks on a base of %s packs a week, measured the way "
                    "the brief measures it. The offer form goes to George Fenelon by 30 October and Tullan "
                    "confirms the leaflet on 13 November." % (thousand(g["packs_Save 40%"]),
                                                              whole(g["base_chain"])), body),
          Image(png, width=160 * mm, height=55.5 * mm)]
    data = [[Paragraph("Multiple of base", cellb), Paragraph("All stores", cellrb)] +
            [Paragraph(r, cellrb) for r in REGIONS]]
    for d in DEALS:
        a, ar = (cellb, cellrb) if d == OFFER else (cell, cellr)
        data.append([Paragraph(d, a), Paragraph(one(g["chain"][d]), ar)] +
                    [Paragraph(one(g["region"][r][d]), ar) for r in REGIONS])
    t = Table(data, colWidths=[30 * mm, 21 * mm, 30 * mm, 20 * mm, 29 * mm, 33 * mm], hAlign="LEFT")
    t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, 0), 0.7, colors.HexColor(NAVY)),
                           ("LINEBELOW", (0, -1), (-1, -1), 0.5, colors.HexColor("#999999")),
                           ("BACKGROUND", (0, 2), (-1, 2), colors.HexColor("#e9eff6")),
                           ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                           ("TOPPADDING", (0, 0), (-1, -1), 1.6), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6)]))
    st += [t, Spacer(1, 1 * mm),
           Paragraph("Forecast for the stores taking part that send scan data; the eleven Appendix B stores are out. "
                     "Regions as they will stand in March: the 6 July 2026 structure, with Phibsborough trading "
                     "as 0388 and Swords Rivervalley (0119) in Dublin &amp; North East. The Waterford and Wexford "
                     "moves take effect in May 2027 and are not applied.", small)]
    st += [Paragraph("Why our own deal history understates March", h2),
           Paragraph("Every Delahunt deal at Tullan has been Save 20%% or Save 25%%, and our reviews put them at %s "
                     "to %s times base. None of them took the 160s near the price of Tullan's own-brand 160s, so "
                     "they cannot tell us what a deeper deal does. Other brands' deals can. Wherever another "
                     "brand's 160s went on deal at or below that store's own-brand 160s price, own-brand 160s sales "
                     "fell by about a third for the three weeks, and the deal line sold extra packs worth %s of "
                     "own brand's normal sales, twice what own brand lost. Where the deal price stayed above own "
                     "brand, neither line moved beyond the usual deal lift." % (one(min(own)), one(max(own)),
                                                                              one(g["share"])), body),
           Paragraph("At March prices Save 40%% takes our 160s to €%.2f in the stores that sell them at €%.2f, "
                     "below the €%.2f those stores charge for own brand. In the stores at €%.2f and €%.2f it stays "
                     "above own brand (€%.2f against €%.2f, €%.2f against €%.2f), so they get the deal lift alone, "
                     "about %s times base. Dublin &amp; North East has most of the €%.2f stores and forecasts %s; "
                     "the other three regions forecast %s to %s. Tullan holds tea prices until the review on "
                     "5 April, so these are the prices the deal runs at."
                     % (d_hi[OFFER], p_hi, o_hi, p_mid, p_lo, d_mid[OFFER], o_mid, d_lo[OFFER], o_lo,
                        one(1 + GD.lift_at(g["pure"], 0.40)), p_hi, one(g["region"]["Dublin & North East"][OFFER]),
                        one(min(rest)), one(max(rest))), body)]
    st += [Paragraph("The other three deals", h2),
           Paragraph("<b>Save a third, %s.</b> At €%.2f, €%.2f and €%.2f the 160s stay above own brand in every "
                     "store, so the deal earns its lift and nothing more." % (ch["Save a third"],
                                                                             d_hi["Save a third"],
                                                                             d_mid["Save a third"],
                                                                             d_lo["Save a third"]), body),
           Paragraph("<b>€2.40 off, %s.</b> €%.2f, €%.2f and €%.2f also stay above own brand everywhere, the "
                     "closest being €%.2f against €%.2f. It only clears if the deal is checked against the "
                     "€3.79 own-brand RRP, and the stores selling our 160s at €%.2f charge €%.2f for own brand, "
                     "not €3.79." % (ch["€2.40 off"], d_hi["€2.40 off"], d_mid["€2.40 off"], d_lo["€2.40 off"],
                                     d_lo["€2.40 off"], o_lo, p_lo, o_lo), body),
           Paragraph("<b>Half price, %s.</b> It clears comfortably, and it is what Eamon expects Tullan to ask "
                     "for, but it costs about twice as much to fund as Save 40%% for a slot Save 40%% already "
                     "wins." % ch["Half price"], body)]
    cl = g["claims"]
    st.append(KeepTogether([
        Paragraph("This year's billing", h2),
        Paragraph("Joanne asked for the 2026 billing to be checked before anyone signs. Tullan over-billed both "
                  "deals: €%s on BD26-04, where Finglas was billed for a cycle it did not run, and €%s on "
                  "BD26-13, where DN-26-1022 re-billed lines DN-26-0947 had already charged and the originals "
                  "were never credited. The workings set out invoiced against supported for each, ex VAT, and we "
                  "will ask Tullan for the credit with the offer." % (whole(cl["BD26-04"][2]),
                                                                       whole(cl["BD26-13"][2])), body),
        Spacer(1, 2 * mm),
        Paragraph("Sources: TullanInsight category feed, w/c 7 October 2024 to w/c 28 September 2026; Big Deals "
                  "Spring 2027 supplier brief (28 September 2026) and its Appendix B; Tullan member list "
                  "(12 January 2026), area notice and 2026 member news; Tullan debit notes, funding lines and "
                  "manual returns; Delahunt deal reviews 2025-26. Store-by-store workings in "
                  "march_offer_workings.xlsx. Tess Tuohey.", small)]))
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=17 * mm, rightMargin=17 * mm, topMargin=12 * mm,
                            bottomMargin=14 * mm, title="March Big Deals: Tullan offer", author="Tess Tuohey",
                            subject="", creator="Revenue Growth", producer="Delahunt Tea")
    doc.build(st, onFirstPage=footer, onLaterPages=footer)


# ---------------------------------------------------------------- the workbook
NORM = Font(name="Calibri", size=10)
BOLD = Font(name="Calibri", bold=True, size=10)
TITLE = Font(name="Calibri", bold=True, size=12)
GREY = Font(name="Calibri", size=9, color="595959", italic=True)
FILL = PatternFill("solid", fgColor="DCE6F1")
PICK = PatternFill("solid", fgColor="FFF2CC")
LINE = Border(top=Side(style="thin", color="7F7F7F"))
EUR = '"€"#,##0'
EUR2 = '"€"0.00'


def put(ws, r, c, v, font=NORM, fmt=None, fill=None, align=None, wrap=False):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if align or wrap:
        cell.alignment = Alignment(horizontal=align, wrap_text=wrap, vertical="center" if wrap else None)
    return cell


def header(ws, r, labels, start=1):
    for j, h in enumerate(labels):
        put(ws, r, start + j, h, BOLD, fill=FILL, align="left" if j == 0 else "right", wrap=True)


def workbook(g, fund, names, notes, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Offer form"
    put(ws, 1, 1, "Big Deals BD27-03, 1 to 21 March 2027: supplier offer, tea", TITLE)
    for i, (k, v) in enumerate([("Supplier", "Delahunt Tea Ltd"), ("Line", "Delahunt Original 160s, item 211402"),
                                ("Mechanic", OFFER), ("Forecast multiple of base", float(one(g["chain"][OFFER])))]):
        put(ws, 3 + i, 1, k, BOLD)
        put(ws, 3 + i, 2, v, fmt="0.0" if isinstance(v, float) else None, align="left")
    header(ws, 8, ["Trading region", "Base, packs a week", "Forecast event packs, three weeks"])
    for i, reg in enumerate(REGIONS):
        put(ws, 9 + i, 1, reg)
        put(ws, 9 + i, 2, int(math.floor(g["base_region"][reg] + 0.5)), fmt="#,##0")
        put(ws, 9 + i, 3, int(math.floor(g["packs_region"][reg] / 1000 + 0.5)) * 1000, fmt="#,##0")
    put(ws, 13, 1, "Total", BOLD).border = LINE
    put(ws, 13, 2, int(math.floor(g["base_chain"] + 0.5)), BOLD, "#,##0").border = LINE
    put(ws, 13, 3, int(math.floor(g["packs_Save 40%"] / 1000 + 0.5)) * 1000, BOLD, "#,##0").border = LINE
    put(ws, 15, 1, "Event packs to the nearest thousand. Definitions on the Notes sheet.", GREY)
    for k, v in {"A": 27, "B": 36, "C": 20}.items():
        ws.column_dimensions[k].width = v
    ws.row_dimensions[8].height = 28

    ws = wb.create_sheet("Forecast by deal")
    put(ws, 1, 1, "Forecast multiple of base, three weeks, by deal and trading region", TITLE)
    header(ws, 3, ["Deal", "All stores"] + REGIONS + ["Clears 2.5"])
    for i, d in enumerate(DEALS):
        f = PICK if d == OFFER else None
        put(ws, 4 + i, 1, d, BOLD if d == OFFER else NORM, fill=f)
        put(ws, 4 + i, 2, float(one(g["chain"][d])), fmt="0.0", fill=f)
        for j, reg in enumerate(REGIONS):
            put(ws, 4 + i, 3 + j, float(one(g["region"][reg][d])), fmt="0.0", fill=f)
        put(ws, 4 + i, 7, "Yes" if g["chain"][d] >= GD.FLOOR else "No", fill=f, align="right")
    put(ws, 9, 1, "Tullan's floor applies to the line across all stores taking part. Store detail on Store workings.",
        GREY)
    for k, v in {"A": 14, "B": 11, "C": 18, "D": 9, "E": 17, "F": 20, "G": 10}.items():
        ws.column_dimensions[k].width = v
    ws.row_dimensions[3].height = 28

    ws = wb.create_sheet("Store workings")
    hdr = ["Store no", "Store", "Trading region, March 2027", "Delahunt 160s price", "Own-brand 160s price",
           "Base, packs a week", "Own-brand base, packs a week"]
    for d in DEALS:
        hdr += [d + " price", d + " at or below own brand", d + " packs a week"]
    header(ws, 1, hdr)
    fc = g["fc"]
    order = sorted(fc)
    for i, s in enumerate(order):
        r = fc[s]
        row = [(s, None), (names[s], None), (g["regions"][s], None), (r["reg"], EUR2), (r["ol_reg"], EUR2),
               (round(r["base"], 2), "0.0"), (round(r["ol_base"], 2), "0.0")]
        for d in DEALS:
            u, dp, cross = r[d]
            row += [(dp, EUR2), ("Yes" if cross else "No", None), (round(u, 2), "0.0")]
        for j, (v, fmt) in enumerate(row):
            put(ws, 2 + i, 1 + j, v, fmt=fmt, align="right" if j == 0 or isinstance(v, str) and j > 6 else None)
    last = 2 + len(order)
    put(ws, last, 1, "Total", BOLD).border = LINE
    for j in range(2, len(hdr) + 1):
        ws.cell(row=last, column=j).border = LINE
    sums = {6: sum(r["base"] for r in fc.values()), 7: sum(r["ol_base"] for r in fc.values())}
    for k, d in enumerate(DEALS):
        sums[10 + 3 * k] = sum(r[d][0] for r in fc.values())
    for col, v in sums.items():
        put(ws, last, col, round(v, 2), BOLD, "#,##0").border = LINE
    ws.freeze_panes = "C2"
    widths = [8, 28, 21, 10, 10, 10, 11] + [10, 11, 10] * 4
    for j, w in enumerate(widths):
        ws.column_dimensions[ws.cell(row=1, column=1 + j).column_letter].width = w
    ws.row_dimensions[1].height = 42
    ws.auto_filter.ref = "A1:%s%d" % (ws.cell(row=1, column=len(hdr)).column_letter, last - 1)
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws = wb.create_sheet("2026 billing check")
    put(ws, 1, 1, "Big Deals funding on Delahunt Original 160s, 2026: what Tullan invoiced against what it should "
                  "have, ex VAT", TITLE)
    header(ws, 3, ["Cycle", "Deal", "Invoiced", "Supported by the deal terms", "Difference"])
    meta = {"BD26-04": "Save 20%, w/c 9 March 2026", "BD26-13": "Save 25%, w/c 14 September 2026"}
    for i, (code, (inv, sup, dif)) in enumerate(sorted(g["claims"].items())):
        put(ws, 4 + i, 1, code)
        put(ws, 4 + i, 2, meta[code])
        put(ws, 4 + i, 3, int(math.floor(inv + 0.5)), fmt=EUR)
        put(ws, 4 + i, 4, int(math.floor(sup + 0.5)), fmt=EUR)
        put(ws, 4 + i, 5, int(math.floor(dif + 0.5)), BOLD, fmt=EUR)
    put(ws, 7, 1, "Debit notes as issued", BOLD)
    header(ws, 8, ["Debit note", "Cycle", "Issued", "Net", "Note"])
    for i, d in enumerate(notes):
        put(ws, 9 + i, 1, d["Debit note"])
        put(ws, 9 + i, 2, d["Cycle"])
        put(ws, 9 + i, 3, d["Issued"], fmt="dd/mm/yyyy", align="right")
        put(ws, 9 + i, 4, round(float(d["Net (EUR)"]), 2), fmt='"€"#,##0.00')
        put(ws, 9 + i, 5, d["Note"] or "")
    r = 9 + len(notes) + 1
    for line in ("BD26-04: Finglas (0256) is billed for the three weeks although it traded through its refit "
                 "without running the cycle; the feed has it at regular price throughout.",
                 "BD26-13: DN-26-1022 re-bills 41 stores' weeks 2 and 3 at the right price and adds three week 1 "
                 "files that came in late. The DN-26-0947 lines it replaces were never credited.",
                 "Supported: regular price less deal price on every pack sold on deal in the scan feed, plus the "
                 "deal packs on the two island stores' monthly returns."):
        put(ws, r, 1, line, GREY)
        r += 1
    for k, v in {"A": 12, "B": 28, "C": 12, "D": 20, "E": 12}.items():
        ws.column_dimensions[k].width = v
    ws.row_dimensions[3].height = 28

    ws = wb.create_sheet("Notes")
    lines = [
        ("Notes", TITLE),
        ("", NORM),
        ("Base: as section 3 of the brief. Each store's average weekly units of the 160s over the weeks "
         "commencing 29 June to 21 September 2026 for which Tullan has the store's file and the line was off "
         "deal, summed over the stores taking part that send scan data.", NORM),
        ("Stores taking part: every scanning member except the eleven in Appendix B. Arranmore Island and Inis Mor send "
         "monthly returns, not scan data, so they are outside the forecast and inside the billing check.", NORM),
        ("Deal prices: each store's regular price less the deal, rounded down to the cent (brief, section 1), at "
         "the prices in the feed for w/c 28 September 2026, which Tullan holds until 5 April 2027.", NORM),
        ("Deal lift: measured on the 2024 to 2026 Big Deals cycles in the feed, each against the line's 13 weeks "
         "before the cycle, in store-weeks where the deal price stayed above the own-brand line of the same bag "
         "count; read between the depths run for €2.40 off.", NORM),
        ("Own-brand switching: in store-weeks where a branded deal price was at or below the store's own-brand "
         "price for the same bag count, own-brand units fell by about a third, and the deal line sold extra packs "
         "equal to 0.7 of the store's own-brand base, read on the deal line as its units above base x (1 + deal "
         "lift). A deal that reaches own brand in a store adds 0.7 of that store's own-brand base.", NORM),
        ("Regions: March 2027 regions, from the January 2026 member list, the 6 July 2026 area changes and the "
         "2026 member news. The May 2027 moves are not applied.", NORM),
        ("Scan feed: TullanInsight category feed v3, w/c 7 October 2024 to w/c 28 September 2026.", NORM),
    ]
    for i, (text, font) in enumerate(lines):
        c = put(ws, 1 + i, 1, text, font)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 110
    wb.properties.creator = "Tess Tuohey"
    wb.properties.lastModifiedBy = "Tess Tuohey"
    wb.save(path)


def main():
    OUT.mkdir(exist_ok=True)
    g, fund, ladder, names, own, notes = figures()
    paper(g, fund, ladder, own, OUT / "march_offer_paper.pdf", chart(g))
    workbook(g, fund, names, notes, OUT / "march_offer_workings.xlsx")
    seal.seal_dir(OUT, seal.GOLDEN)
    return g


if __name__ == "__main__":
    g = main()
    print("pick", OFFER, one(g["chain"][OFFER]), "packs", thousand(g["packs_Save 40%"]), "base", whole(g["base_chain"]))
