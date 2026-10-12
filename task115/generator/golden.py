"""The golden solve: every graded figure, computed from the shipped files in target/ only.

    golden(target)  -> dict of figures

Nothing here imports generator state. The estimators are the plain ones a careful analyst would
write: per-store bases over the 13 weeks before each deal, the deal's lift measured where no
own-label line of the same bag count was undercut, the own-label loss measured where one was."""
import csv
import datetime as dt
import math
import re
from collections import defaultdict
from pathlib import Path

import pandas as pd
from docx import Document
from openpyxl import load_workbook

FEED = "TullanInsight_tea_store_week_2024W41-2026W40.csv"
LINES_F = "tullan_tea_range_2026.csv"
CAL_F = "big_deals_calendar_2024-2026.csv"
MEMBERS_F = "tullan_members_jan2026.xlsx"
AREA_F = "tullan_area_structure_from_6_July_2026.docx"
BULLETIN_F = "member_news_2026.docx"
BRIEF_F = "Big_Deals_Spring_2027_supplier_brief.docx"
CLAIMS_F = "tullan_deal_funding_lines_Delahunt_2026.csv"
DN_F = "tullan_debit_notes_Delahunt_2026.xlsx"
RETURNS_F = "tullan_manual_returns_2026.csv"

BASE_FROM, BASE_TO = dt.date(2026, 6, 29), dt.date(2026, 9, 21)
EVENT_WEEKS = 3
DEL160, OL_BY_BAGS = None, None
MENU = [("Save a third", "pct", 1 / 3), ("Save 40%", "pct", 0.40), ("Half price", "pct", 0.50),
        ("€2.40 off", "abs", 2.40)]
FLOOR = 2.5


def wk(s):
    return dt.date.fromisoformat(s)


def load_feed(T):
    df = pd.read_csv(T / FEED, dtype={"store_no": str, "item_no": str, "deal_ref": str})
    df["week"] = pd.to_datetime(df["week_commencing"]).dt.date
    return df


def load_lines(T):
    out = {}
    with open(T / LINES_F, encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            out[r["item_no"]] = r
    return out


def load_calendar(T):
    ev = []
    with open(T / CAL_F, encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            if r["category"] != "Tea":
                continue
            kind = "pct" if r["deal_type"] == "PCT" else "abs"
            val = float(r["deal_value"]) / (100 if kind == "pct" else 1)
            if kind == "pct" and abs(val - 0.3333) < 1e-3:
                val = 1 / 3
            ev.append({"code": r["cycle"], "start": wk(r["start_week"]), "item": r["item_no"], "kind": kind,
                       "val": val, "label": r["deal_text"]})
    return ev


def deal_price(reg, kind, val):
    x = reg * (1 - val) if kind == "pct" else reg - val
    return math.floor(round(x * 100, 6)) / 100


def doc_tables(path):
    d = Document(str(path))
    out = []
    for t in d.tables:
        rows = [[c.text.strip() for c in r.cells] for r in t.rows]
        out.append(rows)
    return out


def regions_at_event(T):
    """Store number -> trading region in March 2027, from the January list, the July 2026 area changes
    (those already in force) and the year's member news."""
    wb = load_workbook(T / MEMBERS_F, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows[0]]
    i_no, i_reg = hdr.index("Store no"), hdr.index("Trading region")
    reg = {str(r[i_no]).zfill(4): r[i_reg] for r in rows[1:] if r[i_no] is not None}
    for t in doc_tables(T / AREA_F):
        h = t[0]
        if "Store" in h[0] and "From" in h[-1]:
            for r in t[1:]:
                no, frm = r[0].zfill(4), r[-1]
                eff = dt.datetime.strptime(frm, "%d %B %Y").date()
                if eff <= dt.date(2027, 3, 1):
                    reg[no] = r[3]
    for t in doc_tables(T / BULLETIN_F):
        h = t[0]
        if h[:2] == ["Store", "Change"]:
            for r in t[1:]:
                no, change, detail = r[0].zfill(4), r[1], r[2]
                if change == "New store number":
                    new = re.search(r"\b(\d{4})\b", detail).group(1)
                    reg[new] = reg[no]
                elif change == "Opened":
                    m = re.search(r"trading region (.+?)\.?$", detail)
                    reg[no] = m.group(1)
    return reg


def appendix_b(T):
    for t in doc_tables(T / BRIEF_F):
        if t[0][0] == "Store no" and "Reason" in t[0][-1]:
            return {r[0].zfill(4) for r in t[1:]}
    raise ValueError("Appendix B not found")


def store_base(df, item, weeks, exclude_rows=None):
    sub = df[(df.item_no == item) & (df.week.isin(set(weeks))) & (df.deal_price_eur.isna())]
    if exclude_rows is not None:
        sub = sub[~sub.index.isin(exclude_rows)]
    return sub.groupby("store_no").units.mean()


def pure_and_switch(df, lines, cal, window=13, pool="pooled"):
    """Pure lift by depth (pct deals, store-weeks not undercutting own label), then the switch measured on the
    deal line where a deal undercut own label: its units over base x (1 + pure lift at its depth), per own-label
    base. The own-label line's own loss in those store-weeks is returned beside it."""
    bags = {k: int(v["bag_count"]) for k, v in lines.items()}
    own = {bags[k]: k for k, v in lines.items() if v["supplier"] == "Tullan Grocers"}
    by_sw = df.set_index(["store_no", "week", "item_no"])
    num, den = defaultdict(float), defaultdict(float)
    per_event = {}
    ol_num, ol_den = 0.0, 0.0
    gain_num = []
    for e in cal:
        weeks = [e["start"] + dt.timedelta(days=7 * k) for k in range(EVENT_WEEKS)]
        pre = [e["start"] - dt.timedelta(days=7 * k) for k in range(1, window + 1)]
        base = store_base(df, e["item"], pre)
        pair = own.get(bags[e["item"]])
        obase = store_base(df, pair, pre) if pair else None
        ev = df[(df.item_no == e["item"]) & (df.week.isin(set(weeks))) & (df.deal_price_eur.notna())]
        en, ed = 0.0, 0.0
        for r in ev.itertuples():
            if r.store_no not in base.index:
                continue
            cross = False
            if pair:
                try:
                    olp = by_sw.loc[(r.store_no, r.week, pair), "regular_price_eur"]
                    cross = r.deal_price_eur <= float(olp) + 1e-9
                except KeyError:
                    cross = False
            if cross:
                if obase is not None and r.store_no in obase.index:
                    olu = by_sw.loc[(r.store_no, r.week, pair), "units"]
                    ol_num += float(olu)
                    ol_den += float(obase[r.store_no])
                    dep = e["val"] if e["kind"] == "pct" else e["val"] / r.regular_price_eur
                    gain_num.append((dep, float(r.units), float(base[r.store_no]), float(obase[r.store_no])))
                continue
            if e["kind"] != "pct":
                continue
            num[e["val"]] += r.units
            den[e["val"]] += base[r.store_no]
            en += r.units
            ed += base[r.store_no]
        if ed:
            per_event[e["code"]] = (e["val"], en / ed - 1)
    if pool == "pooled":
        pure = {d: num[d] / den[d] - 1 for d in num}
    else:      # mean of events at each depth
        acc = defaultdict(list)
        for code, (d, l) in per_event.items():
            acc[d].append(l)
        pure = {d: sum(v) / len(v) for d, v in acc.items()}
    gain = sum(u - b * (1 + lift_at(pure, dep)) for dep, u, b, o in gain_num) / sum(o for *_, o in gain_num)
    return pure, gain, per_event, 1 - ol_num / ol_den


def lift_at(pure, d):
    pts = sorted([(0.0, 0.0)] + list(pure.items()))
    for (x0, y0), (x1, y1) in zip(pts, pts[1:]):
        if d <= x1 + 1e-12:
            return y0 + (y1 - y0) * (d - x0) / (x1 - x0)
    (x0, y0), (x1, y1) = pts[-2], pts[-1]
    return y1 + (y1 - y0) * (d - x1) / (x1 - x0)


def forecast(df, lines, pure, share, stores, test="store", rrp=None):
    """Weekly event units and base for every menu deal, per store."""
    del160 = next(k for k, v in lines.items() if v["description"] == "DELAHUNT ORIGINAL 160S")
    ol160 = next(k for k, v in lines.items() if v["supplier"] == "Tullan Grocers" and v["bag_count"] == "160")
    weeks = [w for w in sorted(df.week.unique()) if BASE_FROM <= w <= BASE_TO]
    b = store_base(df, del160, weeks)
    o = store_base(df, ol160, weeks)
    last = df.week.max()
    cur = df[df.week == last].set_index(["store_no", "item_no"])["regular_price_eur"]
    out = {}
    for s in stores:
        if s not in b.index:
            continue
        reg = float(cur[(s, del160)])
        olp = float(cur[(s, ol160)])
        row = {"base": float(b[s]), "ol_base": float(o.get(s, 0.0)), "reg": reg, "ol_reg": olp}
        for name, kind, val in MENU:
            dp = deal_price(reg, kind, val)
            dpt = val if kind == "pct" else val / reg
            ref = olp if test == "store" else rrp
            cross = dp <= ref + 1e-9
            row[name] = (row["base"] * (1 + lift_at(pure, dpt)) + (share * row["ol_base"] if cross else 0.0),
                         dp, cross)
        out[s] = row
    return out


def summarise(fc, regions):
    res = {"chain": {}, "region": defaultdict(dict), "base_region": defaultdict(float), "packs_region": {}}
    B = sum(r["base"] for r in fc.values())
    res["base_chain"] = B
    for name, _, _ in MENU:
        ev = sum(r[name][0] for r in fc.values())
        res["chain"][name] = ev / B
        res["packs_" + name] = EVENT_WEEKS * ev
    for s, r in fc.items():
        res["base_region"][regions[s]] += r["base"]
    for reg in res["base_region"]:
        ss = [s for s in fc if regions[s] == reg]
        Br = sum(fc[s]["base"] for s in ss)
        for name, _, _ in MENU:
            res["region"][reg][name] = sum(fc[s][name][0] for s in ss) / Br
        res["packs_region"][reg] = EVENT_WEEKS * sum(fc[s]["Save 40%"][0] for s in ss)
    return res


def claims_check(T, df):
    """Invoiced against supported, ex VAT, for each of Delahunt's 2026 deals."""
    wb = load_workbook(T / DN_F, read_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    hdr = [str(h) for h in rows[0]]
    dn = [dict(zip(hdr, r)) for r in rows[1:] if r[0]]
    invoiced = defaultdict(float)
    for d in dn:
        invoiced[d["Cycle"]] += float(d["Net (EUR)"])
    lines = pd.read_csv(T / CLAIMS_F, dtype={"store_no": str, "item_no": str})
    ret = pd.read_csv(T / RETURNS_F, dtype={"store_no": str, "item_no": str})
    supported = defaultdict(float)
    del160 = "211402"
    for code in invoiced:
        ev = df[(df.deal_ref == code) & (df.item_no == del160)]
        supported[code] += float(((ev.regular_price_eur - ev.deal_price_eur) * ev.units).sum())
        rr = ret[(ret.deal_ref == code) & (ret.item_no == del160)]
        supported[code] += float(((rr.regular_price_eur - rr.deal_price_eur) * rr.deal_units).sum())
    return {c: (round(invoiced[c], 2), round(supported[c], 2), round(invoiced[c] - supported[c], 2))
            for c in sorted(invoiced)}


def golden(T):
    T = Path(T)
    df = load_feed(T)
    lines = load_lines(T)
    cal = load_calendar(T)
    pure, share, per_event, loss = pure_and_switch(df, lines, cal)
    regions = regions_at_event(T)
    out_b = appendix_b(T)
    weeks = [w for w in sorted(df.week.unique()) if BASE_FROM <= w <= BASE_TO]
    in_window = set(df[df.week.isin(set(weeks))].store_no)
    stores = sorted(s for s in in_window if s not in out_b)
    fc = forecast(df, lines, pure, share, stores)
    res = summarise(fc, regions)
    res["pure"] = pure
    res["share"] = share
    res["loss_share"] = loss
    res["per_event"] = per_event
    res["claims"] = claims_check(T, df)
    res["stores"] = stores
    res["fc"] = fc
    res["regions"] = {s: regions[s] for s in stores}
    return res
