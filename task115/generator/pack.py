"""Write the data files of the bundle (CSV and XLSX). Documents are in docs.py."""
import csv
import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

import world as W

FEED = "TullanInsight_tea_store_week_2024W41-2026W40.csv"


def write_feed(rows, T):
    with open(T / FEED, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["week_commencing", "store_no", "item_no", "units", "sales_eur", "regular_price_eur",
                    "deal_price_eur", "deal_ref"])
        for (wk, sid, item, u, reg, dp, ref) in rows:
            paid = dp if dp is not None else reg
            w.writerow([wk.isoformat(), sid, item, u, "%.2f" % (u * paid), "%.2f" % reg,
                        "" if dp is None else "%.2f" % dp, ref or ""])
    return len(rows)


def write_lines(T):
    with open(T / "tullan_tea_range_2026.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["item_no", "gtin", "supplier", "description", "bag_count", "case_size", "express_range",
                    "status"])
        for item, l in W.LINE.items():
            w.writerow([item, l["gtin"], l["supplier"], l["description"] if "description" in l else l["desc"],
                        l["bags"], l["case"], "Y" if l["express"] else "N", "Active"])


OTHER_FEATURES = [
    ("Coffee", "402215", "BELLAMONT ROAST & GROUND 227G", "PCT", 25), ("Biscuits", "512033", "KILCAR DIGESTIVES 400G",
                                                                        "PCT", 33.33),
    ("Cereal", "611470", "ARDFIN PORRIDGE OATS 1KG", "MONEY", 1.00), ("Coffee", "402219", "BELLAMONT INSTANT 200G",
                                                                       "PCT", 40),
    ("Biscuits", "512041", "KILCAR CHOC CHIP COOKIES 200G", "PCT", 50), ("Bakery", "705112", "HILLVIEW BROWN SODA 500G",
                                                                        "MONEY", 0.50),
    ("Cereal", "611488", "ARDFIN GRANOLA 750G", "PCT", 25), ("Coffee", "402230", "BELLAMONT PODS 10S", "PCT", 33.33),
]


def write_calendar(T):
    import itertools
    rows = []
    for k in range(0, 40):
        start = W.cycle_start(k)
        if start < dt.date(2024, 10, 21) or start > dt.date(2026, 9, 14):
            continue
        code = None
        for e in W.EVENTS:
            if e[1] == start:
                code = e[0]
        rows.append(start)
    rows = sorted(set(rows))
    out = []
    feats = itertools.cycle(OTHER_FEATURES)
    for start in rows:
        yr = start.year
        grid = [W.cycle_start(k) for k in range(0, 60)]
        idx = sum(1 for s in grid if s.year == yr and s <= start)
        code = "BD%02d-%02d" % (yr % 100, idx)
        end = start + dt.timedelta(days=20)
        for _ in range(2):
            cat, item, desc, typ, val = next(feats)
            out.append([code, start.isoformat(), end.isoformat(), cat, item, desc, typ,
                        ("%.2f" % val) if typ == "MONEY" else ("%g" % val),
                        ("€%.2f off" % val) if typ == "MONEY" else ({33.33: "Save a third", 50: "Half price"}.get(
                            val, "Save %g%%" % val))])
        for e in W.EVENTS:
            if e[1] == start:
                assert e[0] == code, (e[0], code)
                typ = "PCT" if e[3] == "pct" else "MONEY"
                val = (e[4] * 100) if e[3] == "pct" else e[4]
                vtxt = "33.33" if abs(val - 33.3333) < 0.01 else ("%.2f" % val if typ == "MONEY" else "%g" % val)
                out.append([code, start.isoformat(), end.isoformat(), "Tea", e[2], W.LINE[e[2]]["desc"], typ, vtxt,
                            e[5]])
    out.sort(key=lambda r: (r[1], r[3] != "Tea", r[4]))
    with open(T / "big_deals_calendar_2024-2026.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["cycle", "start_week", "end_week", "category", "item_no", "description", "deal_type",
                    "deal_value", "deal_text"])
        w.writerows(out)
    return out


def _sheet(ws, header, rows, widths=None):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE4EE")
    for r in rows:
        ws.append(r)
    if widths:
        for col, wdt in widths.items():
            ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "A2"


def write_members(stores, T):
    wb = Workbook()
    ws = wb.active
    ws.title = "Members Jan 2026"
    rows = []
    for s in sorted(stores, key=lambda x: x["no"]):
        if s.get("reissued"):
            continue          # opened August 2026, after this list
        if s.get("closed") and s["closed"] < dt.date(2026, 1, 12):
            continue          # Ballyfermot closed 31 December 2025
        rows.append([s["no"], s["name"], s["town"].replace(" Island", "") if False else s["town"], s["county"],
                     s["routing"], s["format"], W.REGIONS[["DNE", "S", "WNW", "MSE"].index(s["reg_jan"])],
                     s["opened"].strftime("%d/%m/%Y") if s["opened"] > dt.date(2009, 1, 1) else "",
                     "Y" if s["scanning"] else "N"])
    _sheet(ws, ["Store no", "Store name", "Town", "County", "Eircode routing key", "Format", "Trading region",
                "Member since", "Scan data"], rows,
           {"A": 9, "B": 34, "C": 24, "D": 12, "E": 10, "F": 13, "G": 24, "H": 13, "I": 10})
    ws2 = wb.create_sheet("About")
    ws2["A1"] = "Tullan Grocers member list, extracted 12 January 2026 by Member Services."
    ws2["A2"] = "Trading regions as at the extract date. Member since is shown for stores that joined after 2008."
    wb.save(T / "tullan_members_jan2026.xlsx")
    return len(rows)


def write_claims(lines, dn, returns, T):
    with open(T / "tullan_deal_funding_lines_Delahunt_2026.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(lines[0].keys()))
        w.writeheader()
        for l in lines:
            w.writerow(l)
    wb = Workbook()
    ws = wb.active
    ws.title = "Debit notes"
    hdr = list(dn[0].keys())
    _sheet(ws, hdr, [[d[h] if not isinstance(d[h], dt.date) else d[h] for h in hdr] for d in dn],
           {"A": 13, "B": 9, "C": 12, "D": 62, "E": 12, "F": 14, "G": 13, "H": 9, "I": 70, "J": 14})
    for r in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        for c in r:
            c.number_format = "dd/mm/yyyy"
    wb.save(T / "tullan_debit_notes_Delahunt_2026.xlsx")
    with open(T / "tullan_manual_returns_2026.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(returns[0].keys()))
        w.writeheader()
        for r in returns:
            w.writerow(r)


def write_reviews(reviews, T):
    wb = Workbook()
    ws = wb.active
    ws.title = "Deal reviews"
    _sheet(ws, ["Cycle", "Line", "Deal", "Weeks on deal", "Base weeks (Tullan)", "Base packs a week",
                "Deal packs a week", "Multiple of base", "Funding invoiced net (EUR)", "Reviewed by"],
           reviews, {"A": 9, "B": 26, "C": 12, "D": 22, "E": 26, "F": 17, "G": 17, "H": 16, "I": 24, "J": 18})
    ws2 = wb.create_sheet("Notes")
    ws2["A1"] = "Post-deal reviews for Delahunt lines in Tullan Big Deals. Base and deal packs as Tullan's review " \
                "reports them; funding from Tullan's debit notes."
    wb.save(T / "delahunt_deal_reviews_2025-26.xlsx")
