"""The pack's documents: short, functional, in the voice of the people who wrote them."""
import datetime as dt

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab import rl_config
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import seal
import world as W

rl_config.invariant = 1          # fixed dates and /ID, so the bytes rebuild identically; seal.py dates them
REG = dict(zip(["DNE", "S", "WNW", "MSE"], W.REGIONS))


# ---------------------------------------------------------------- helpers
def _pdf(path, blocks, title):
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=9.2, leading=12.2,
                          alignment=TA_LEFT, spaceAfter=4)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=13, spaceAfter=5)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10, spaceBefore=6,
                        spaceAfter=3)
    small = ParagraphStyle("s", parent=body, fontSize=7.8, textColor=colors.HexColor("#4a4a4a"))
    cell = ParagraphStyle("c", fontName="Helvetica", fontSize=8, leading=9.6)
    story = []
    for kind, val in blocks:
        if kind == "h1":
            story.append(Paragraph(val, h1))
        elif kind == "h2":
            story.append(Paragraph(val, h2))
        elif kind == "p":
            story.append(Paragraph(val, body))
        elif kind == "small":
            story.append(Paragraph(val, small))
        elif kind == "sp":
            story.append(Spacer(1, val * mm))
        elif kind == "table":
            data, widths = val
            data = [[Paragraph(str(c), cell) for c in r] for r in data]
            t = Table(data, colWidths=[w_ * mm for w_ in widths] if widths else None, hAlign="LEFT",
                      repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#9a9a9a")),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e4e8ee")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("TOPPADDING", (0, 0), (-1, -1), 1.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5)]))
            story.append(t)
    author, dept, _ = seal.TARGET[path.name]
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=16 * mm, title=title, author=author, creator=dept, producer=author,
                            subject="")
    doc.build(story)


def _docx(path, build):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    build(d)
    d.save(str(path))


def _table(d, rows, bold_header=True):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j)
            c.text = str(v)
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    if i == 0 and bold_header:
                        run.bold = True
    return t


# ---------------------------------------------------------------- Tullan documents
def brief(path, appendix):
    def build(d):
        p = d.add_paragraph()
        r = p.add_run("Tullan Grocers  |  Buying, Hot Beverages")
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        d.add_heading("Big Deals Spring 2027: supplier brief, tea", level=1)
        d.add_paragraph("Cycle BD27-03, Monday 1 March to Sunday 21 March 2027. Issued 28 September 2026.")
        d.add_paragraph("Spring Big Deals is our biggest leaflet of the first half. Tea has one slot. Featured lines "
                        "get the front-of-leaflet tea panel and a gondola end in every participating store for the "
                        "three weeks.")
        d.add_heading("1. Mechanics", level=2)
        d.add_paragraph("For tea we are offering four mechanics: Save a third, Save 40%, Half price, and €2.40 off. "
                        "A deal price is the store's regular shelf price less the deal, rounded down to the cent. "
                        "Regular prices on tea lines are held until the price review on 5 April 2027.")
        d.add_heading("2. Acceptance", level=2)
        d.add_paragraph("We take a line into the event only if the supplier's forecast for the three weeks, across "
                        "participating stores that send us scan data, averages at least two and a half times the "
                        "line's base a week. After the event we check what the line delivered against the same "
                        "base. A line that falls short is not offered a slot in the following two Big Deals.")
        d.add_heading("3. Base", level=2)
        d.add_paragraph("A line's base is measured over the 13 complete trading weeks before the date of this brief, "
                        "weeks commencing 29 June to 21 September 2026. For each participating store we take the "
                        "average weekly units over the weeks in that period for which we received the store's file "
                        "and the line was not on a Big Deal. The line's base is the sum of those store averages over "
                        "the participating stores that send scan data.")
        d.add_heading("4. Funding", level=2)
        d.add_paragraph("The supplier funds the shopper's saving on every unit sold on deal in participating stores, "
                        "at each store's regular price less its deal price. We invoice the funding by debit note "
                        "after the event, from the scan data, and from the monthly returns for members that do not "
                        "send scan data. Where a debit note is re-issued, the re-issue replaces the original's lines "
                        "for the stores and weeks it covers. Debit notes carry VAT at 23%.")
        d.add_heading("5. Offers", level=2)
        d.add_paragraph("Offers on the attached form by 5pm on Friday 30 October 2026, to George Fenelon: the line, "
                        "the mechanic, the line's base in packs a week, the forecast event packs for each trading "
                        "region and in total, and the forecast multiple of base. We confirm the leaflet on "
                        "13 November 2026.")
        d.add_heading("6. Participating stores", level=2)
        d.add_paragraph("All member stores take part except those in Appendix B.")
        d.add_heading("Appendix B. Stores not taking part in BD27-03", level=2)
        _table(d, [["Store no", "Store", "Reason"]] + appendix)
        q = d.add_paragraph("Tullan Grocers Ltd. Commercial in confidence. For the named supplier only.")
        q.runs[0].font.size = Pt(8)
    _docx(path, build)


def area_notice(path, moves):
    def build(d):
        d.add_heading("Trading regions from 6 July 2026", level=1)
        d.add_paragraph("Member Services, 22 June 2026. To all members, area managers and supplier contacts.")
        d.add_paragraph("We are rebalancing the four trading regions so each area manager has a workable number of "
                        "stores. The Kildare stores move to Dublin & North East, and the Limerick and Clare stores "
                        "move to South, from Monday 6 July 2026. The new structure applies to all trading, reporting "
                        "and supplier forecasting from that date. The second stage, moving the Waterford and Wexford "
                        "stores to South, takes effect on Monday 3 May 2027. Area manager contacts are on the "
                        "members' portal.")
        _table(d, [["Store no", "Store", "From region", "To region", "From"]] + moves)
        d.add_paragraph("Questions to Martina Mulgrew, Member Services.")
    _docx(path, build)


def member_news(path, changes):
    def build(d):
        d.add_heading("Member news 2026: store changes", level=1)
        d.add_paragraph("Extracts from the April, June and August issues of Member News. Full issues are on the "
                        "members' portal.")
        d.add_heading("April", level=2)
        d.add_paragraph("Welcome to the new owners of Tullan Phibsborough, who took over on 6 April. The store "
                        "trades under a new store number from that date; nothing changes for customers.")
        d.add_heading("June", level=2)
        d.add_paragraph("Tullan Ballinrobe leaves the group on 31 May after eighteen years. Our thanks to the "
                        "family for everything they gave the group. The summer charity bake sale raised just over "
                        "€41,000 across 160 stores.")
        d.add_heading("August", level=2)
        d.add_paragraph("Tullan Swords Rivervalley opened its doors on 3 August, our second store in Swords. "
                        "Store numbers are reused once a number has been free for two months, so the new store "
                        "takes 0119.")
        d.add_heading("Store changes", level=2)
        _table(d, [["Store", "Change", "Detail"]] + changes)
    _docx(path, build)


def review_pdf(path, rows, notes):
    data = [["Cycle", "Weeks", "Line", "Deal", "Stores", "Base packs a week", "Deal packs a week",
             "Multiple of base"]] + rows
    blocks = [
        ("small", "Tullan Grocers &nbsp;|&nbsp; Category Insights &nbsp;|&nbsp; Big Deals review, tea &nbsp;|&nbsp; "
                  "Prepared 6 October 2026 for supplier reviews"),
        ("h1", "Big Deals review: tea, January 2025 to September 2026"),
        ("p", "Every tea line featured in a Big Deals cycle from BD25-01 to BD26-13. For each cycle, the line's base "
              "is the average weekly units over the eight complete weeks ending two weeks before the cycle "
              "started, leaving out weeks the line was on deal, summed over the stores that ran the cycle. Deal "
              "packs a week is the line's average weekly units over the three cycle weeks in those stores. The "
              "multiples describe the cycles listed; they are not forecasts for any future cycle."),
        ("table", (data, [15, 34, 46, 20, 12, 18, 18, 17])),
        ("sp", 3),
        ("h2", "Notes"),
    ] + [("p", n) for n in notes] + [
        ("p", "Store counts are the stores that ran the cycle and sent scan data for all three weeks."),
        ("small", "Commercial in confidence. Shared with tea suppliers for their own cycle reviews."),
    ]
    _pdf(path, blocks, "Big Deals review: tea 2025-26")


def rrp_pdf(path):
    data = [["Item", "Line", "Recommended price"],
            ["300812", "Tullan Irish Breakfast 80s", "€2.05"], ["300813", "Tullan Irish Breakfast 160s", "€3.79"],
            ["300814", "Tullan Irish Breakfast 240s", "€5.45"], ["300870", "Tullan Decaf 80s", "€2.69"],
            ["301402", "Tullan Ground Coffee 227g", "€3.49"], ["301455", "Tullan Instant Coffee 200g", "€4.99"],
            ["310021", "Tullan Digestives 400g", "€1.19"], ["310040", "Tullan Porridge Oats 1kg", "€1.69"],
            ["320118", "Tullan Irish Butter 454g", "€4.29"], ["320230", "Tullan Free Range Eggs 6", "€2.39"]]
    blocks = [
        ("small", "Tullan Grocers &nbsp;|&nbsp; Own Brand"),
        ("h1", "Tullan own brand: recommended retail prices from 5 January 2026"),
        ("p", "Recommended retail prices for the grocery lines in the January 2026 own brand review. Members set "
              "their own shelf prices."),
        ("table", (data, [22, 80, 34])),
        ("sp", 2),
        ("small", "Own Brand team, 15 December 2025."),
    ]
    _pdf(path, blocks, "Tullan own brand RRP 2026")


def feed_spec_pdf(path):
    data = [["Field", "Meaning"],
            ["week_commencing", "Monday of the trading week (Monday to Sunday)."],
            ["store_no", "Four-digit member store number. A store that changes number reports under the new number "
                         "from the week of the change."],
            ["item_no", "Tullan item number (see the range file for GTIN, supplier and bag count)."],
            ["units", "Selling units scanned in the week. Every item ranged in the store has a row, zero where "
                      "nothing sold."],
            ["sales_eur", "Value scanned at the price paid, including VAT where it applies."],
            ["regular_price_eur", "The store's regular shelf price for the item in that week."],
            ["deal_price_eur", "The Big Deals price, filled only where the item was on a Big Deal in that store "
                               "and week."],
            ["deal_ref", "The Big Deals cycle, filled only where deal_price_eur is."]]
    blocks = [
        ("small", "Tullan Grocers &nbsp;|&nbsp; Data Services"),
        ("h1", "TullanInsight category feed, version 3: field reference"),
        ("p", "The category feed gives suppliers in a category the scan data for every line in it, store by "
              "week. Files are cut on Monday morning for the week to the previous Sunday."),
        ("table", (data, [36, 130])),
        ("sp", 2),
        ("p", "One row per store, week and item. A store-week with no rows means we did not receive the store's "
              "file for that week; it does not mean the store sold nothing. Members that do not send scan data "
              "report through the monthly returns instead and are not in this feed."),
        ("small", "Version 3, March 2024. Data Services, Tullan Grocers."),
    ]
    _pdf(path, blocks, "TullanInsight category feed v3")


# ---------------------------------------------------------------- Delahunt documents
def thread(path):
    txt = """From: Joanne Uniacke <juniacke@delahunttea.ie>
Sent: Wednesday 7 October 2026 16:42
To: Tess Tuohey; Eamon Manahan
Cc: Rebekah Masterson
Subject: RE: Tullan spring Big Deals - offer

Tess,

Before I put my name to anything for March I want this year's Tullan billing looked at properly.
We paid what came in and nobody checked it against anything.

Joanne

-----Original Message-----
From: Eamon Manahan <emanahan@delahunttea.ie>
Sent: Wednesday 7 October 2026 11:05
To: Tess Tuohey
Cc: Joanne Uniacke; Rebekah Masterson
Subject: RE: Tullan spring Big Deals - offer

Tess, I was with George on Monday. My read is they want half price for the spring leaflet and
anything less won't get the slot. I'd rather we went in strong than not at all.

Eamon

-----Original Message-----
From: Rebekah Masterson <rmasterson@delahunttea.ie>
Sent: Tuesday 6 October 2026 17:20
To: Tess Tuohey
Cc: Eamon Manahan; Joanne Uniacke
Subject: RE: Tullan spring Big Deals - offer

Our own deals at Tullan have only ever been Save 20% and Save 25% on the 160s, and the reviews
are in the shared folder. The new category feed came through this morning, two years of it,
every line in tea, so you have Carrow, Carnmore and Rathmore as well as ours.

Rebekah

-----Original Message-----
From: Tess Tuohey <ttuohey@delahunttea.ie>
Sent: Tuesday 6 October 2026 09:12
To: Rebekah Masterson
Cc: Eamon Manahan; Joanne Uniacke
Subject: Tullan spring Big Deals - offer

Rebekah, can you pull everything we have from Tullan for the spring brief? Offer is due on the 30th,
leadership sees it on the 22nd.

Thanks, Tess

This email and any attachments are confidential. Delahunt Tea Ltd, Little Island, Cork.
"""
    path.write_text(txt, encoding="utf-8")


def readme(path, counts):
    txt = """TULLAN SPRING 2027 - FILES RECEIVED AND PULLED
Delahunt Tea, Revenue Growth. Folder made up 8 October 2026 (R. Masterson).

From Tullan Grocers (supplier portal and buyer emails)
  TullanInsight_tea_store_week_2024W41-2026W40.csv   category feed, all tea lines, store by week,
                                                     w/c 7 Oct 2024 to w/c 28 Sep 2026, {feed} rows
  TullanInsight_category_feed_v3_spec.pdf            Tullan's field reference for the feed
  tullan_tea_range_2026.csv                          Tullan's tea range file (items, GTINs, bag counts)
  big_deals_calendar_2024-2026.csv                   Big Deals cycles from Oct 2024, all categories
  big_deals_review_tea_2025-26.pdf                   Tullan Category Insights review of tea cycles
  Big_Deals_Spring_2027_supplier_brief.docx          the spring brief, issued 28 Sep 2026
  tullan_members_jan2026.xlsx                        member list as at 12 Jan 2026 ({members} stores)
  tullan_area_structure_from_6_July_2026.docx        Member Services notice on trading regions
  member_news_2026.docx                              store changes from Member News, 2026
  tullan_own_brand_rrp_2026.pdf                      own brand recommended prices, Jan 2026
  tullan_debit_notes_Delahunt_2026.xlsx              debit notes for our 2026 Big Deals funding
  tullan_deal_funding_lines_Delahunt_2026.csv        the lines behind those debit notes ({claims} lines)
  tullan_manual_returns_2026.csv                     monthly returns from the two members without scan data

Delahunt
  delahunt_deal_reviews_2025-26.xlsx                 our post-deal reviews of Tullan cycles
  delahunt_commercial_thread_oct2026.txt             email thread on the spring offer

All Tullan files are shared under our supplier agreement and are commercial in confidence.
Store routing keys are the Eircode routing keys Tullan holds for each store.
"""
    path.write_text(txt.format(**counts), encoding="utf-8")
