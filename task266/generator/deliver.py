"""The goldens for task266: program_completion_review_2027.xlsx (each program's three-year completion share and whether
it enters review, its fall 2026 FTE and its classes under ten students, the completion curves, the notes) and
completion_review_board_paper_2027.pdf (the board paper with the recommendation and the chart), plus
generator/figures.json and the task's metadata.json. Every figure comes from analysis.py run on the shipped pack."""
import json
import os
import shutil
import tempfile

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt                          # noqa: E402
import numpy as np                                       # noqa: E402
from openpyxl import Workbook                            # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

import analysis as A                                     # noqa: E402
import common as K                                       # noqa: E402
import docs as D                                         # noqa: E402
import postprocess as PP                                 # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
DELIVERABLES = ["program_completion_review_2027.xlsx", "completion_review_board_paper_2027.pdf"]
INV = {v: k for k, v in K.SIS.items()}


def r1(x):
    return float("%.1f" % (np.floor(x * 10 + 0.5) / 10))


def results(target):
    R, pk, ent = A.call(target)
    progs = [K.SIS[p] for p in K.CODES]
    inrev = [p for p in progs if R[p]["golden"] < K.FLOOR]
    a = A.ask_a(target)
    b = A.ask_b(target)
    stops_a = dict(natural=A.ask_a(target, credits="row", program="row", census="status"),
                   row_credits=A.ask_a(target, credits="row"),
                   row_credits_once=A.ask_a(target, credits="row_once"),
                   extract_program=A.ask_a(target, program="row"),
                   late_adds_kept=A.ask_a(target, census="status"),
                   withdrawals_dropped=A.ask_a(target, census="all_dates"))
    stops_b = dict(natural=A.ask_b(target, enrol="export", group="none"),
                   census_uncombined=A.ask_b(target, group="none"),
                   meeting=A.ask_b(target, group="meeting"),
                   text=A.ask_b(target, group="text"),
                   room=A.ask_b(target, group="room"),
                   late_adds_kept=A.ask_b(target, enrol="status", group="meeting"))
    secs = A.rows_xlsx(os.path.join(target, K.FILES["schedule"]), "Fall 2026 sections")
    regs = A.rows_csv(os.path.join(target, K.FILES["registrations"]))
    n = {}
    for r in A.census_rows(regs, A.day(pk.census[K.FALL26]), "date"):
        n[r["crn"]] = n.get(r["crn"], 0) + 1
    nur_clinical = sum(1 for s in secs if s["Subj"] == "NUR" and s["Type"] == "CLN" and s["Status"] == "A"
                       and n.get(str(s["CRN"]), 0) < K.SMALL)
    return dict(R=R, progs=progs, inrev=inrev, a=a, b=b, stops_a=stops_a, stops_b=stops_b, ent=ent,
                nur_clinical=nur_clinical, pk=pk)


# ---------------------------------------------------------------------------------------------- the scorecard
def _head(ws, head, widths, fill="22465F"):
    ws.append(head)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.row_dimensions[ws.max_row].height = 44


def write_scorecard(path, res):
    R = res["R"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Scorecard"
    ws.append(["Program review 2027: three-year completion and the fall 2026 profile"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Entering students of fall 2023, 2024 and 2025 followed to the end of spring 2027; review floor 30.0 per "
               "cent (Board Policy 3.12, 3.4)."])
    ws["A2"].font = Font(italic=True, color="5B646B")
    ws.append([None])
    _head(ws, ["Program", "Award", "Program code", "Entering students, fall 2023 to 2025",
               "Three-year completion share (%)", "Enters review", "FTE at the fall 2026 census",
               "Classes under ten students, fall 2026"], [44, 7, 12, 14, 14, 9, 13, 14])
    ws.freeze_panes = "B5"
    first = ws.max_row + 1
    for p in res["progs"]:
        code = INV[p]
        ws.append([K.NAME[code], K.AWARD[code], p, R[p]["n"], r1(R[p]["golden"]),
                   "Yes" if p in res["inrev"] else "No", r1(res["a"][p]), res["b"][p]])
        row = ws.max_row
        ws.cell(row, 4).number_format = "#,##0"
        ws.cell(row, 5).number_format = "0.0"
        ws.cell(row, 7).number_format = "#,##0.0"
        for c in (4, 5, 7, 8):
            ws.cell(row, c).alignment = Alignment(horizontal="right")
        ws.cell(row, 6).alignment = Alignment(horizontal="center")
        if p in res["inrev"]:
            ws.cell(row, 6).font = Font(bold=True, color="9C2B1E")
    last = ws.max_row
    ws.append(["All nine programs", None, None, "=SUM(D%d:D%d)" % (first, last), None,
               "=COUNTIF(F%d:F%d,\"Yes\")" % (first, last), "=SUM(G%d:G%d)" % (first, last),
               "=SUM(H%d:H%d)" % (first, last)])
    row = ws.max_row
    for c in range(1, 9):
        ws.cell(row, c).font = Font(bold=True)
    ws.cell(row, 4).number_format = "#,##0"
    ws.cell(row, 7).number_format = "#,##0.0"
    ws.append([None])
    ws.append(["Students who died (%d) are removed from the classes. FTE and classes are the review profile of Policy "
               "3.12, section 4." % sum(R[p]["deaths"] for p in res["progs"])])
    ws.cell(ws.max_row, 1).font = Font(italic=True, color="5B646B")
    ws.append(["%d of Nursing's %d classes under ten are clinical groups, which take at most eight students each." % (
        res["nur_clinical"], res["b"][K.SIS["NUR"]])])
    ws.cell(ws.max_row, 1).font = Font(italic=True, color="5B646B")
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1

    ws2 = wb.create_sheet("Completion by term")
    _head(ws2, ["Term since entry"] + [K.SHORTNAME[INV[p]] for p in res["progs"]], [12] + [11] * 9)
    ws2.freeze_panes = "B2"
    for t in range(K.WINDOW):
        ws2.append([t + 1] + [r1(R[p]["curve"][t]) for p in res["progs"]])
        for c in range(2, 11):
            ws2.cell(ws2.max_row, c).number_format = "0.0"
    ws2.append([None])
    ws2.append(["Share of the program's entering students who had earned an associate degree by the end of each term, "
                "per cent; term 1 is the fall term of entry and summer terms count."])
    ws2.cell(ws2.max_row, 1).font = Font(italic=True, color="5B646B")

    ws3 = wb.create_sheet("Notes")
    ws3.column_dimensions["A"].width = 118
    notes = [
        "Notes",
        None,
        "Entering students: the Office of Institutional Research's fall entering-student files for fall 2023, 2024 "
        "and 2025 (fall 2023 on its legacy student numbers, carried to the current system through the 2024 "
        "conversion map), each student in the program entered.",
        "Outcome: an associate degree from the college, in any program, under the admission recorded at entry, by the "
        "end of the ninth term counting the fall of entry (Policy 3.12, 3.1 and 3.2).",
        "Each admission's ending and its term come from the registrar's admissions extract of 28 May 2027: a degree "
        "(GR), a transfer out (TO) or a withdrawal (WD); admissions with no exit were open at the end of spring 2027.",
        "Estimate: for each program, term by term, the share of the admissions still open at the start of the term "
        "that ended in a degree in that term, times the share of the class still open; transfers out and withdrawals "
        "end an admission without the degree, and open admissions count until the end of spring 2027 (the "
        "cumulative incidence of the degree, with transfer out and withdrawal as the other endings).",
        "The fall 2024 class is observed to its eighth term and the fall 2025 class to its fifth; their later terms "
        "use the rates of the classes observed in those terms.",
        "FTE: credit hours registered at the fall 2026 census (registrations in place on 8 September 2026, later "
        "withdrawals included and registrations added after that date not), each at its section's credit hours in "
        "the fall 2026 schedule (laboratory, shop and clinical sections carry none, and English Composition I and "
        "College Algebra ran at three), by each student's program of record for fall 2026, divided by 15.",
        "Classes under ten students: sections of each program's own subjects that ran in fall 2026, sections taught "
        "together at one meeting (interactive video between campuses) counted as one class with their students "
        "added; the accelerated sections that share a room and time in the first and second eight weeks are two "
        "classes; census registrations as above.",
        "Shares are rounded to one decimal place from the unrounded estimate; the floor is applied to the unrounded "
        "share.",
        "Prepared by %s, Director of Institutional Research, 4 June 2027." % K.PEOPLE["ir"],
    ]
    for line in notes:
        ws3.append([line])
        ws3.cell(ws3.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws3["A1"].font = Font(bold=True, size=12)
    wb.save(path)
    PP.COMPANY = K.COLLEGE
    PP.normalize_ooxml(path, K.PEOPLE["ir"], K.PEOPLE["ir"], "2027-06-01T13:12:40Z", "2027-06-04T15:47:18Z",
                       excel=True)


# ---------------------------------------------------------------------------------------------- the chart
def chart(path, res):
    R = res["R"]
    fig, ax = plt.subplots(figsize=(7.2, 3.9), dpi=200)
    x = np.arange(0, K.WINDOW + 1)
    rev_col = ["#9c2b1e", "#c4512f", "#d9822b", "#b5654a", "#7d3c2f", "#e0a05a"]
    ok_col = ["#22465f", "#3d6fa3", "#5c8a64"]
    i_r = i_o = 0
    ends = []
    for p in res["progs"]:
        y = [0.0] + list(R[p]["curve"])
        if p in res["inrev"]:
            col, i_r = rev_col[i_r], i_r + 1
            ls = "-"
        else:
            col, i_o = ok_col[i_o], i_o + 1
            ls = "-"
        ax.step(x, y, where="post", color=col, lw=1.6 if p in res["inrev"] else 1.9, ls=ls)
        ends.append([y[-1], K.SHORTNAME[INV[p]], col])
    ax.axhline(K.FLOOR, color="#1d2326", lw=1.1, ls=(0, (2, 2)))
    ax.text(0.15, K.FLOOR + .7, "Review floor, 30.0%", fontsize=7.5, color="#1d2326")
    ends.sort()
    placed = []
    for v, name, col in ends:
        yv = v
        for q in placed:
            if abs(yv - q) < 1.9:
                yv = q + 1.9
        placed.append(yv)
        ax.text(9.12, yv, "%s %.1f" % (name, r1(v)), fontsize=6.6, color=col, va="center")
    ax.set_xlim(0, 11.6)
    ax.set_ylim(0, 46)
    ax.set_xticks(range(0, 10))
    ax.set_xlabel("Terms since entry (the fall of entry is term 1; summer terms count)", fontsize=8)
    ax.set_ylabel("Entering students with an associate degree (%)", fontsize=8)
    ax.tick_params(labelsize=7.5)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.set_title("Six programs end their three years below the floor", fontsize=9.5, loc="left", color="#1d2326")
    fig.tight_layout()
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)


# ---------------------------------------------------------------------------------------------- the board paper
def write_paper(path, res, png):
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, Spacer
    R = res["R"]
    S = D.styles("Sans", colors.HexColor("#22465f"))
    doc = D._doc(path, "Academic program review 2027: three-year completion", K.COLLEGE)
    names = [K.NAME[INV[p]] for p in res["inrev"]]
    keep = [K.NAME[INV[p]] for p in res["progs"] if p not in res["inrev"]]
    st = [Paragraph("Six of nine associate-degree programs enter review in 2027", S["h1"]),
          Paragraph("Board of Trustees, regular meeting of 17 June 2027, agenda item 7.2: Academic program review "
                    "(Board Policy 3.12). From %s, Director of Institutional Research; through %s, Vice President "
                    "for Academic Affairs. 4 June 2027." % (K.PEOPLE["ir"], K.PEOPLE["vpaa"]), S["meta"])]
    P = lambda t: st.append(Paragraph(t, S["body"]))
    st.append(Paragraph("Recommendation", S["h2"]))
    rev = [K.NAME[INV[p]] for p in res["inrev"]]
    shares = ["%.1f" % r1(R[p]["golden"]) for p in res["inrev"]]
    out = [(K.SHORTNAME[INV[p]], "%.1f" % r1(R[p]["golden"])) for p in res["progs"] if p not in res["inrev"]]
    P("That the board place six programs in review from 2027-28: %s and %s. Their three-year completion shares, %s "
      "and %s per cent, are below the policy's floor of 30 per cent of entering students earning an associate degree "
      "within three years. %s (%s per cent), %s (%s) and %s (%s) stay out of review." % (
          ", ".join(rev[:-1]), rev[-1], ", ".join(shares[:-1]), shares[-1], out[0][0], out[0][1], out[1][0],
          out[1][1], out[2][0], out[2][1]))
    rows = [["Program", "Entering students", "Three-year completion (%)", "Review", "FTE, fall 2026",
             "Classes under ten"]]
    for p in res["progs"]:
        rows.append([K.NAME[INV[p]] if INV[p] != "HVR" else "HVAC-R Technology", "{:,}".format(R[p]["n"]),
                     "%.1f" % r1(R[p]["golden"]), "Yes" if p in res["inrev"] else "No",
                     "{:,.1f}".format(r1(res["a"][p])), str(res["b"][p])])
    st.append(D._table(S, rows, [62, 24, 30, 15, 24, 22], right=(1, 2, 4, 5)))
    st.append(Paragraph("Sources: IR fall entering-student files; registrar's admissions extract, conferrals, "
                        "fall 2026 registrations, schedule and program history, 28 May 2027.", S["meta"]))
    st.append(Image(png, width=172 * mm, height=172 * mm * 3.9 / 7.2))
    st.append(Spacer(1, 4))
    st.append(Paragraph("How the share is measured", S["h2"]))
    P("The policy counts every student who entered in fall 2023, 2024 or 2025, including the two classes still inside "
      "their three years. Each admission ends with a degree, a transfer to another institution or a withdrawal, or "
      "is still open. For each program we take, term by term, the share of the admissions still open that ended in a "
      "degree in that term, weighted by the share of the class still open; a student who transferred out or "
      "withdrew has ended the admission without the college's degree, and only the admissions still open at the end "
      "of spring 2027 are carried forward on the rates of the classes observed in the later terms. Treating "
      "transfers and withdrawals as students whose outcome is still unknown would credit each of them with a share "
      "of a degree they can no longer earn here: it puts every program except Early Childhood Education above the "
      "floor, the transfer-oriented programs at %s per cent." % (
          " to ".join("%.0f" % v for v in (min(R[p]["op"] for p in res["progs"] if INV[p] in K.TRANSFER),
                                            max(R[p]["op"] for p in res["progs"] if INV[p] in K.TRANSFER)))))
    P("The fall 2023 class, the only one whose three years have closed, is not the measure: the guided-pathways "
      "advising begun in fall 2024 has cut first-year withdrawals, and on that class alone Criminal Justice "
      "(%.1f per cent) would also enter review." % r1(R[K.SIS["CJ"]]["closed2023"]))
    st.append(Paragraph("Review profile", S["h2"]))
    P("FTE counts the credit hours registered at the fall 2026 census by each program's students, every section at "
      "the credit hours it carried in fall 2026, divided by 15. Classes under ten counts the program's own sections "
      "that ran in fall 2026 with fewer than ten students at the census, sections taught together over interactive "
      "video counted once and the accelerated eight-week sections counted separately. %d of Nursing's %d are "
      "clinical groups, which take at most eight students each." % (res["nur_clinical"], res["b"][K.SIS["NUR"]]))
    st.append(Paragraph("Notes. Students who died (%d) are removed from the classes. Shares are rounded to one "
                        "decimal place; the floor applies to the unrounded share. The workbook "
                        "program_completion_review_2027.xlsx has each program's share by term." % sum(
                            R[p]["deaths"] for p in res["progs"]), S["meta"]))
    doc.build(st, onFirstPage=D._frame("Board of Trustees, 17 June 2027", "Agenda item 7.2", "Sans"),
              onLaterPages=D._frame("Board of Trustees, 17 June 2027", "Agenda item 7.2", "Sans"))
    PP.normalize_pdf(path, b"D:20270604101522-04'00'", b"D:20270604154831-04'00'")


# ---------------------------------------------------------------------------------------------- figures.json
def figures(res):
    R = res["R"]
    rung_keys = ["golden", "op", "op_exits_first", "crude", "complete", "closed2023", "half_wd", "half_to", "strat",
                 "person", "redate", "dc_compete", "dc_censor", "final", "own"]
    return dict(
        call=dict(in_review=[K.NAME[INV[p]] for p in res["inrev"]], count=len(res["inrev"]),
                  share={K.NAME[INV[p]]: r1(R[p]["golden"]) for p in res["progs"]},
                  unrounded={K.NAME[INV[p]]: R[p]["golden"] for p in res["progs"]}),
        entrants={K.NAME[INV[p]]: R[p]["n"] for p in res["progs"]},
        deaths=sum(R[p]["deaths"] for p in res["progs"]),
        rungs={k: {K.NAME[INV[p]]: R[p][k] for p in res["progs"]} for k in rung_keys},
        rung_lists={k: [K.NAME[INV[p]] for p in res["progs"] if R[p][k] < K.FLOOR] for k in rung_keys},
        curves={K.NAME[INV[p]]: [r1(v) for v in R[p]["curve"]] for p in res["progs"]},
        ask_a={K.NAME[INV[p]]: r1(res["a"][p]) for p in res["progs"]},
        ask_a_unrounded={K.NAME[INV[p]]: res["a"][p] for p in res["progs"]},
        ask_a_stops={k: {K.NAME[INV[p]]: r1(v[p]) for p in res["progs"]} for k, v in res["stops_a"].items()},
        ask_b={K.NAME[INV[p]]: res["b"][p] for p in res["progs"]},
        ask_b_stops={k: {K.NAME[INV[p]]: v[p] for p in res["progs"]} for k, v in res["stops_b"].items()},
        floor=K.FLOOR, fte_hours=K.FTE_HOURS, small=K.SMALL, window_terms=K.WINDOW,
        entering_in_files=len(res["ent"]), entering_total=sum(R[p]["n"] for p in res["progs"]),
        observed_terms={str(c): res["pk"].idx(K.fall(c), K.LAST) for c in K.CLASSES},
        nur_clinical=res["nur_clinical"],
        rival_text={"op_transfer_low": int(np.floor(min(R[p]["op"] for p in res["progs"] if INV[p] in K.TRANSFER) + .5)),
                    "op_transfer_high": int(np.floor(max(R[p]["op"] for p in res["progs"] if INV[p] in K.TRANSFER) + .5)),
                    "cj_2023": r1(R[K.SIS["CJ"]]["closed2023"])},
    )


def write(target, golden, T):
    res = results(target)
    if os.path.exists(golden):
        shutil.rmtree(golden)
    os.makedirs(golden)
    write_scorecard(os.path.join(golden, DELIVERABLES[0]), res)
    tmp = tempfile.mkdtemp()
    png = os.path.join(tmp, "completion_chart.png")
    chart(png, res)
    write_paper(os.path.join(golden, DELIVERABLES[1]), res, png)
    shutil.rmtree(tmp)
    F = figures(res)
    json.dump(F, open(os.path.join(HERE, "figures.json"), "w"), indent=1, sort_keys=True)
    meta = {"distractor_files": ["target/" + K.FILES[k] for k in K.DISTRACTORS],
            "deliverables": DELIVERABLES,
            "stated_inputs": ["30", "2023", "2024", "2025", "2026", "2027"],
            "as_of": "2027-05-28"}
    json.dump(meta, open(os.path.join(os.path.dirname(HERE), "metadata.json"), "w"), indent=1)
    T["res"] = res
    return res
