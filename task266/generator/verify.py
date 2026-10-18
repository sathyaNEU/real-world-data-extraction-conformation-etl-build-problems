"""The independent verifier: reads only the shipped bundle (python3 verify.py <target>), recomputes the call, the
rungs and every graded ask figure on its own code path (the csv module and openpyxl, an estimator written from the
counts of each term's endings among the admissions still open, no generator state, no analysis.py), and compares
everything with generator/figures.json. Exit 0 when every figure agrees."""
import csv
import json
import os
import re
import sys
from collections import Counter, defaultdict

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
FAIL = []
FLOOR, NINE, SMALL, FTE = 30.0, 9, 10, 15


def same(label, a, b):
    if a != b:
        FAIL.append("%s: verifier %r, figures.json %r" % (label, a, b))


def rows_csv(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def sheet(path, name):
    ws = load_workbook(path, read_only=True, data_only=True)[name]
    it = ws.iter_rows(values_only=True)
    head = next(it)
    return [dict(zip(head, r)) for r in it if any(v is not None for v in r)]


def round1(x):
    return int(x * 10 + 0.5) / 10


def incidence(rec, compete=True):
    """rec: (term index of the ending or None, ending, last term observed). The award's cumulative incidence over
    nine terms, the other endings competing (compete) or censored at their term (not compete)."""
    left = {}
    for t in range(1, NINE + 1):
        left[t] = sum(1 for e, k, o in rec if (e is not None and e >= t) or (e is None and o >= t))
    surv, cum = 1.0, 0.0
    for t in range(1, NINE + 1):
        n = left[t]
        if not n:
            continue
        g = sum(1 for e, k, o in rec if e == t and k == "GR")
        x = sum(1 for e, k, o in rec if e == t and k != "GR")
        if compete:
            cum += surv * g / n
            surv *= 1 - (g + x) / n
        else:
            surv *= 1 - g / n
    return 100 * (cum if compete else 1 - surv)


def main(target):
    P = lambda name: os.path.join(target, name)
    files = os.listdir(target)
    find = lambda pat: [f for f in files if re.search(pat, f)][0]
    F = json.load(open(os.path.join(HERE, "figures.json")))

    terms = sheet(P(find(r"term_calendar")), "Terms")
    order = [str(t["Term code"]) for t in terms]
    legacy = {t["Legacy term code"]: str(t["Term code"]) for t in terms if t["Legacy term code"]}
    progs = sheet(P(find(r"program_register")), "Programs")
    name = {r["Program code"]: r["Program"] for r in progs}
    leg_prog = {r["Legacy code"]: r["Program code"] for r in progs}
    nine = [r["Program code"] for r in progs if r["Award"] in ("AA", "AS", "AAS")]
    owner = {r["Subject"]: r["Program code"] for r in sheet(P(find(r"program_register")), "Subjects")}
    newid = {}
    for line in open(P(find(r"conversion")), encoding="utf-8").read().splitlines()[1:]:
        old, new = line.split("|")[:2]
        newid[old] = new

    # ------------------------------------------------------------------ the entering students and their endings
    wb = P(find(r"entering_students"))
    ent = [(newid[str(r["ID"])], legacy[r["Term"]], leg_prog[r["Prog"]]) for r in sheet(wb, "Fall 2023")]
    for y in (2024, 2025):
        ent += [(r["Student ID"], str(r["Cohort Term"]), r["Program Code"]) for r in sheet(wb, "Fall %d" % y)]
    ending = {(a["student_id"], a["admit_term"]): (a["exit_term"], a["exit_reason"])
              for a in rows_csv(P(find(r"admissions_exits")))}
    last = "202720"
    rec = defaultdict(list)
    first23 = defaultdict(list)
    for sid, admit, prog in ent:
        code, why = ending[(sid, admit)]
        if why == "DC":
            continue
        obs = order.index(last) - order.index(admit) + 1
        if not why:
            r = (None, "C", obs)
        else:
            r = (order.index(code) - order.index(admit) + 1, why, obs)
        rec[prog].append(r)
        if admit == "202410":
            first23[prog].append(r)
    share = {name[p]: round1(incidence(rec[p])) for p in nine}
    same("the nine shares", share, F["call"]["share"])
    same("the review list", [name[p] for p in nine if incidence(rec[p]) < FLOOR], F["call"]["in_review"])
    same("entering students", {name[p]: len(rec[p]) for p in nine}, F["entrants"])
    km = {name[p]: incidence(rec[p], compete=False) for p in nine}
    same("the default path's list", [n for n, v in km.items() if v < FLOOR], F["rung_lists"]["op"])
    for n in km:
        if abs(km[n] - F["rungs"]["op"][n]) > 1e-9:
            FAIL.append("default path %s: %r against %r" % (n, km[n], F["rungs"]["op"][n]))
    closed = {name[p]: 100 * sum(1 for e, k, o in first23[p] if k == "GR" and e <= NINE) / len(first23[p]) for p in nine}
    same("the 2023 class list", [n for n, v in closed.items() if v < FLOOR], F["rung_lists"]["closed2023"])

    # ------------------------------------------------------------------ Ask A: FTE at the fall 2026 census
    import datetime as _dt
    cal = sheet(P(find(r"term_calendar")), "Terms")
    census = [t["Census date"] for t in cal if str(t["Term code"]) == "202710"][0]
    census = census.date() if hasattr(census, "date") else census

    def when(v):
        return _dt.date(int(v[:4]), int(v[5:7]), int(v[8:10]))

    secs = sheet(P(find(r"section_schedule")), "Fall 2026 sections")
    cr = {str(s["CRN"]): s["Cr Hrs"] for s in secs}
    hist = rows_csv(P(find(r"program_of_record")))
    at = {}
    for h in hist:
        if h["effective_term"] <= "202710" and (not h["end_term"] or h["end_term"] >= "202710"):
            at[h["student_id"]] = h["program_code"]
    regs = rows_csv(P(find(r"census_registrations")))
    # in place on the census date: registered by then, or withdrawn after it (a W grade)
    held = [r for r in regs if r["reg_status"] == "WC" or
            (r["reg_status"] in ("RE", "RW") and when(r["status_date"]) <= census)]
    hours = Counter()
    for r in held:
        hours[at.get(r["student_id"])] += cr[r["crn"]]
    same("Ask A", {name[p]: round1(hours[p] / FTE) for p in nine}, F["ask_a"])

    # ------------------------------------------------------------------ Ask B: classes under ten
    n = Counter(r["crn"] for r in held)
    meet = defaultdict(list)
    for s in secs:
        if s["Status"] != "A" or owner.get(s["Subj"]) not in nine:
            continue
        key = str(s["CRN"]) if s["Days"] == "TBA" else (s["Instructor"], s["Days"], s["Begin"])
        meet[key].append(s)
    small = Counter()
    for lst in meet.values():
        # sections at one instructor's meeting time are one class only over the weeks they share
        lst = sorted(lst, key=lambda s: (s["Start Date"], s["CRN"]))
        runs = [[lst[0]]]
        for s in lst[1:]:
            if s["Start Date"] <= max(x["End Date"] for x in runs[-1]):
                runs[-1].append(s)
            else:
                runs.append([s])
        for run in runs:
            if sum(n[str(s["CRN"])] for s in run) < SMALL:
                small[owner[run[0]["Subj"]]] += 1
    same("Ask B", {name[p]: small[p] for p in nine}, F["ask_b"])
    if FAIL:
        print("verify: %d mismatches" % len(FAIL))
        for x in FAIL:
            print("  " + x)
        sys.exit(1)
    print("verify: call, rungs and every ask figure agree with figures.json (%d entering students, %d registrations)"
          % (len(ent), len(regs)))


if __name__ == "__main__":
    main(sys.argv[1])
