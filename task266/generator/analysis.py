"""Every figure from the shipped pack alone: the main call (each program's three-year completion share, the award's
cumulative incidence with transfer out and withdrawal as competing endings and open admissions censored), every rung
of the ladder and every variant the fork table asserts, and both asks with their stops."""
import csv
import datetime as dt
import os
from collections import Counter, defaultdict

from openpyxl import load_workbook

import common as K

FS = ("FA", "SP")


def rows_csv(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def rows_xlsx(path, sheet):
    ws = load_workbook(path, read_only=True, data_only=True)[sheet]
    it = ws.iter_rows(values_only=True)
    head = next(it)
    return [dict(zip(head, r)) for r in it if any(v is not None for v in r)]


class Pack:
    def __init__(self, target):
        P = lambda k: os.path.join(target, K.FILES[k])
        self.P = P
        terms = rows_xlsx(P("terms"), "Terms")
        self.order = [str(t["Term code"]) for t in terms]
        self.legacy_term = {t["Legacy term code"]: str(t["Term code"]) for t in terms if t["Legacy term code"]}
        self.season = {str(t["Term code"]): str(t["Term"]).split()[0] for t in terms}
        self.census = {str(t["Term code"]): t["Census date"] for t in terms}
        self.confer = {str(t["Term code"]): t["Conferral date"] for t in terms}
        self.start = {str(t["Term code"]): t["First day"] for t in terms}
        reg = rows_xlsx(P("register"), "Programs")
        self.legacy_prog = {r["Legacy code"]: r["Program code"] for r in reg}
        self.programs = [r["Program code"] for r in reg if r["Award"] in ("AA", "AS", "AAS")]
        self.owner = {r["Subject"]: r["Program code"] for r in rows_xlsx(P("register"), "Subjects")}
        self.idmap = {}
        with open(P("idmap"), encoding="utf-8") as fh:
            next(fh)
            for line in fh:
                lid, sid, _, _ = line.rstrip("\r\n").split("|")
                self.idmap[lid] = sid
        self.adm = rows_csv(P("admissions"))
        self.by_adm = {(a["student_id"], a["admit_term"]): a for a in self.adm}
        self.later = defaultdict(list)
        for a in self.adm:
            self.later[a["student_id"]].append(a)

    def idx(self, entry, code):
        return self.order.index(code) - self.order.index(entry) + 1

    def is_fs(self, code):
        return self.season[code] in ("Fall", "Spring")


# ---------------------------------------------------------------------------------------------- the entering classes
def entering(pk):
    """IR's three fall entering files on the current system's numbers, codes and terms."""
    out = []
    for r in rows_xlsx(pk.P("cohorts"), "Fall 2023"):
        out.append(dict(sid=pk.idmap[str(r["ID"])], admit=pk.legacy_term[r["Term"]], prog=pk.legacy_prog[r["Prog"]],
                        cls=2023))
    for cls in (2024, 2025):
        for r in rows_xlsx(pk.P("cohorts"), "Fall %d" % cls):
            out.append(dict(sid=r["Student ID"], admit=str(r["Cohort Term"]), prog=r["Program Code"], cls=cls))
    return out


def spells(pk, ent, mode="golden"):
    """(event index or None, kind, last observed index) per entering student: kind GR, XTO, XWD (another ending) or
    C. mode: golden; person (followed across a readmission); redate (a transfer or withdrawal one fall or spring term
    earlier); dc_compete; dc_censor."""
    later = pk.later
    out = []
    for e in ent:
        a = pk.by_adm[(e["sid"], e["admit"])]
        obs = pk.idx(e["admit"], K.LAST)
        reason, code = a["exit_reason"], a["exit_term"]
        if mode == "person" and reason == "WD":
            nxt = sorted((x for x in later[e["sid"]] if x["admit_term"] > code and x["admit_type"] == "R"),
                         key=lambda x: x["admit_term"])
            if nxt:
                reason, code = nxt[0]["exit_reason"], nxt[0]["exit_term"]
        if reason == "DC":
            if mode == "dc_compete":
                out.append((pk.idx(e["admit"], code), "XDC", obs))
            elif mode == "dc_censor":
                out.append((None, "C", pk.idx(e["admit"], code)))
            continue
        if not reason:
            out.append((None, "C", obs))
            continue
        t = pk.idx(e["admit"], code)
        if mode == "redate" and reason in ("TO", "WD"):
            back = [c for c in pk.order if e["admit"] <= c < code and pk.is_fs(c)]
            t = pk.idx(e["admit"], back[-1]) if back else t
        out.append((t, "GR" if reason == "GR" else "X" + reason, obs))
    return out


def at_risk(sp, t):
    return sum(1 for e, k, o in sp if (e is not None and e >= t) or (e is None and o >= t))


def aj(sp, horizon=K.WINDOW, compete=("XTO", "XWD", "XDC"), curve=False):
    S, F, path = 1.0, 0.0, []
    for t in range(1, horizon + 1):
        n = at_risk(sp, t)
        a = sum(1 for e, k, o in sp if e == t and k == "GR")
        d = sum(1 for e, k, o in sp if e == t and (k == "GR" or k in compete))
        if n:
            F += S * a / n
            S *= 1 - d / n
        path.append(100 * F)
    return path if curve else 100 * F


def km(sp, horizon=K.WINDOW, exits_first=False, curve=False):
    S, path = 1.0, []
    for t in range(1, horizon + 1):
        n = at_risk(sp, t)
        a = sum(1 for e, k, o in sp if e == t and k == "GR")
        if exits_first:
            n -= sum(1 for e, k, o in sp if e == t and k != "GR")
        if n:
            S *= 1 - a / n
        path.append(100 * (1 - S))
    return path if curve else 100 * (1 - S)


def crude(sp, horizon=K.WINDOW):
    return 100 * sum(1 for e, k, o in sp if k == "GR" and e <= horizon) / len(sp)


def complete_case(sp, horizon=K.WINDOW):
    closed = [(e, k) for e, k, o in sp if e is not None and e <= min(o, horizon)]
    return 100 * sum(1 for e, k in closed if k == "GR") / len(closed)


def stratified(pk, ent, prog):
    """Class by class, each class's own rates where it was observed and the observed classes' pooled rates beyond,
    weighted by class size."""
    by = {c: spells(pk, [e for e in ent if e["prog"] == prog and e["cls"] == c]) for c in K.CLASSES}
    obs = {c: pk.idx(K.fall(c), K.LAST) for c in K.CLASSES}
    num = tot = 0
    for c in K.CLASSES:
        S, F = 1.0, 0.0
        for t in range(1, K.WINDOW + 1):
            pool = by[c] if obs[c] >= t else [x for cc in K.CLASSES if obs[cc] >= t for x in by[cc]]
            n = at_risk(pool, t)
            a = sum(1 for e, k, o in pool if e == t and k == "GR")
            d = sum(1 for e, k, o in pool if e == t)
            if n:
                F += S * a / n
                S *= 1 - d / n
        num += F * len(by[c])
        tot += len(by[c])
    return 100 * num / tot


def by_conferral_program(pk, ent):
    """Each entering student placed in the program of the degree, where the admission ended in one."""
    conf = defaultdict(list)
    for c in rows_csv(pk.P("conferrals")):
        conf[c["student_id"]].append(c)
    out = []
    for e in ent:
        a = pk.by_adm[(e["sid"], e["admit"])]
        prog = e["prog"]
        if a["exit_reason"] == "GR":
            cs = [c for c in conf[e["sid"]] if c["term_code"] == a["exit_term"]]
            prog = cs[0]["program_code"]
        out.append(dict(e, prog=prog))
    return out


def own_degree_only(pk, ent):
    """A degree in another program than the one entered counted as another ending."""
    conf = defaultdict(list)
    for c in rows_csv(pk.P("conferrals")):
        conf[c["student_id"]].append(c)
    other = set()
    for e in ent:
        a = pk.by_adm[(e["sid"], e["admit"])]
        if a["exit_reason"] == "GR":
            cs = [c for c in conf[e["sid"]] if c["term_code"] == a["exit_term"]]
            if cs[0]["program_code"] != e["prog"]:
                other.add(e["sid"])
    return other


def call(target):
    """Every rung and variant per program; the golden is 'golden'."""
    pk = Pack(target)
    ent = entering(pk)
    fin = by_conferral_program(pk, ent)
    other = own_degree_only(pk, ent)
    R = {}
    for p in pk.programs:
        es = [e for e in ent if e["prog"] == p]
        sp = spells(pk, es)
        own = []
        for e in es:
            for ev, k, o in spells(pk, [e]):
                own.append((ev, "XWD" if (k == "GR" and e["sid"] in other) else k, o))
        R[p] = dict(
            golden=aj(sp), op=km(sp), op_exits_first=km(sp, exits_first=True), crude=crude(sp),
            complete=complete_case(sp), closed2023=crude(spells(pk, [e for e in es if e["cls"] == 2023])),
            half_wd=aj(sp, compete=("XWD",)), half_to=aj(sp, compete=("XTO",)), strat=stratified(pk, ent, p),
            person=aj(spells(pk, es, "person")), redate=aj(spells(pk, es, "redate")),
            dc_compete=aj(spells(pk, es, "dc_compete")), dc_censor=aj(spells(pk, es, "dc_censor")),
            final=aj(spells(pk, [e for e in fin if e["prog"] == p])), own=aj(own),
            curve=aj(sp, curve=True), curve_op=km(sp, curve=True),
            n=len(sp), entrants=len(es), deaths=len(es) - len(sp),
            awarded=sum(1 for e, k, o in sp if k == "GR" and e <= K.WINDOW),
            open=sum(1 for e, k, o in sp if k == "C"),
            classes={c: sum(1 for e in es if e["cls"] == c) for c in K.CLASSES})
    return R, pk, ent


def review(R, rung="golden"):
    return [p for p in R if R[p][rung] < K.FLOOR]


# ---------------------------------------------------------------------------------------------- the asks
def fall26(target):
    pk = Pack(target)
    regs = rows_csv(pk.P("registrations"))
    secs = rows_xlsx(pk.P("schedule"), "Fall 2026 sections")
    hist = rows_csv(pk.P("history"))
    return pk, regs, secs, hist


def program_at_census(hist, term=K.FALL26):
    out = {}
    for h in hist:
        if h["effective_term"] <= term and (not h["end_term"] or h["end_term"] >= term):
            out[h["student_id"]] = h["program_code"]
    return out


def day(v):
    return v.date() if isinstance(v, dt.datetime) else (dt.date.fromisoformat(v) if isinstance(v, str) else v)


def census_rows(regs, census=None, mode="date"):
    """Registrations in place at the census. mode: 'date' (registered on or before the census date, or withdrawn
    after it); 'status' (RE, RW and WC whatever the date); 'all_dates' (RE, RW and WC dated on or before the census)."""
    out = []
    for r in regs:
        st = r["reg_status"]
        if st not in ("RE", "RW", "WC"):
            continue
        if mode == "date" and st in ("RE", "RW") and day(r["status_date"]) > census:
            continue
        if mode == "all_dates" and day(r["status_date"]) > census:
            continue
        out.append(r)
    return out


def ask_a(target, credits="section", program="census", census="date"):
    """Each program's FTE at the fall 2026 census. credits: 'section' (the schedule's section credit hours), 'row'
    (the registration row's credit_hours) or 'row_once' (the row's credit_hours once per student and course);
    program: 'census' (the history's row in effect for fall 2026) or 'row' (the registration's program_code);
    census: as census_rows."""
    pk, regs, secs, hist = fall26(target)
    sec_cr = {str(s["CRN"]): s["Cr Hrs"] for s in secs}
    prog_at = program_at_census(hist)
    tot = Counter()
    seen = set()
    for r in census_rows(regs, day(pk.census[K.FALL26]), census):
        p = prog_at.get(r["student_id"]) if program == "census" else r["program_code"]
        if credits == "section":
            c = sec_cr[r["crn"]]
        elif credits == "row":
            c = int(r["credit_hours"])
        else:
            key = (r["student_id"], r["subject"], r["course_number"])
            c = 0 if key in seen else int(r["credit_hours"])
            seen.add(key)
        tot[p] += c
    return {p: tot[p] / K.FTE_HOURS for p in pk.programs}


def ask_b(target, enrol="census", group="dated"):
    """Each program's classes with fewer than ten students at the census. enrol: 'census' (registrations in place on
    the census date), 'status' (RE, RW and WC whatever the date) or 'export' (the schedule's Enrl); group: 'dated'
    (sections sharing an instructor and a scheduled meeting in weeks they have in common, whatever the room), 'meeting'
    (sharing an instructor, days and begin time, whatever the weeks), 'exact' (sharing an instructor, days, begin time
    and both dates), 'first' (sharing an instructor, days, begin time and first day), 'none', 'room' (sharing a room,
    days and begin time) or 'text' (instructor, days and begin time as written, TBA included)."""
    pk, regs, secs, hist = fall26(target)
    cen = day(pk.census[K.FALL26])
    n = Counter(r["crn"] for r in census_rows(regs, cen, "date" if enrol == "census" else "status"))
    run = [s for s in secs if s["Status"] == "A" and pk.owner.get(s["Subj"]) in pk.programs]

    def key(s):
        if group == "none":
            return str(s["CRN"])
        if group == "room":
            if s["Room"] in ("ONLINE", "OFF CAMPUS") or s["Days"] == "TBA":
                return str(s["CRN"])
            return (s["Room"], s["Days"], s["Begin"])
        if group == "text":
            return (s["Instructor"], s["Days"], s["Begin"])
        if s["Days"] == "TBA":
            return str(s["CRN"])
        if group == "exact":
            return (s["Instructor"], s["Days"], s["Begin"], day(s["Start Date"]), day(s["End Date"]))
        if group == "first":
            return (s["Instructor"], s["Days"], s["Begin"], day(s["Start Date"]))
        return (s["Instructor"], s["Days"], s["Begin"])

    groups = defaultdict(list)
    for s in run:
        groups[key(s)].append(s)
    if group == "dated":
        split = []
        for lst in groups.values():
            lst = sorted(lst, key=lambda s: (day(s["Start Date"]), s["CRN"]))
            cur, last = [], None
            for s in lst:
                if cur and day(s["Start Date"]) > last:
                    split.append(cur)
                    cur, last = [], None
                cur.append(s)
                last = max(last, day(s["End Date"])) if last else day(s["End Date"])
            split.append(cur)
        classes = split
    else:
        classes = list(groups.values())
    out = Counter()
    for lst in classes:
        size = sum((n[str(s["CRN"])] if enrol != "export" else s["Enrl"]) for s in lst)
        progs = {pk.owner[s["Subj"]] for s in lst}
        if size < K.SMALL:
            for p in progs:
                out[p] += 1
    return {p: out[p] for p in pk.programs}
