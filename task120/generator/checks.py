"""Every assertion the build makes: the ladder's acceptance criteria and the grid, the semantic fork table (SF1 to
SF13), the closure rows, the ask devices and their compositions, the separation line, the input gates, the
counter-pin sweep and the realism sweep."""
import csv
import datetime as dt
import math
import os
import re
import zipfile

import numpy as np
from openpyxl import load_workbook

import asks as A
import analysis as AN
import core
import docs as D
import pack as P

N = [0]


def ok(cond, msg):
    assert cond, msg
    N[0] += 1


def docx_text(path):
    with zipfile.ZipFile(path) as z:
        return re.sub(r"<[^>]+>", " ", z.read("word/document.xml").decode("utf-8"))


def xlsx_strings(path):
    wb = load_workbook(path, read_only=True)
    return "\n".join(str(v) for ws in wb.worksheets for r in ws.iter_rows(values_only=True) for v in r
                     if isinstance(v, str))


def ladder(W, F):
    gap = F["gap"]
    frac = gap - math.floor(gap)
    ok(F["gold_day"] == dt.date(2026, 8, 20), "golden day %s" % F["gold_day"])
    ok(4.45e6 <= F["gold_use"] <= 4.49e6, "golden use %d" % F["gold_use"])
    ok(0.55 <= frac <= 0.95, "gap fraction %.3f" % frac)
    ok(F["call"] == F["call_round"] == math.ceil(gap), "rounding up and rounding off agree (SF7)")
    ok(F["R0"]["figure"] <= F["firm_gal"] - 57600, "R0's production under firm by 40 gpm")
    ok(F["R0"]["call"] is None and F["R0"]["day"] != F["gold_day"], "R0 files none on another day")
    ok(F["gold_prod_rank"] >= 3, "golden day ranks %d on production" % F["gold_prod_rank"])
    ok(F["second_ratio"] <= 0.95, "second use day %.3f of the maximum" % F["second_ratio"])
    for k, v in F["grid"].items():
        if k == ("all", "local"):
            ok(v["call"] == F["call"] and v["day"] == F["gold_day"], "grid golden cell")
            continue
        ok(v["call"] is None or abs(v["call"] - F["call"]) >= 11, "grid cell %s too near: %s" % (k, v["call"]))
    for r in ("R1", "R2", "R3"):
        ok(F[r]["call"] is None or abs(F[r]["call"] - F["call"]) >= 11, "%s too near" % r)
    ok(F["R1"]["call"] != F["R2"]["call"], "R1 and R2 file different figures")
    Bal = W["Bal"]
    ms = [k for k, d in enumerate(Bal["days"]) if d.day == 1] + [len(Bal["days"])]
    mdS = [int(Bal["dS"][a:b].sum()) for a, b in zip(ms[:-1], ms[1:])]
    ok(min(abs(x) for x in mdS) >= 5000, "every month's storage change at least 5,000 gal: %s" % mdS)
    ok(abs(int(Bal["dS"].sum())) >= 5000, "the year's storage change")
    # use equals production less storage change, day by day, and production is whole gallons
    ok(all(int(u) == int(p) - int(s) for u, p, s in zip(Bal["U"], Bal["P"], Bal["dS"])), "balance identity")


def forks(W, F, target):
    H, Bal = W["H"], W["Bal"]
    letter = D.letter_text()
    fac = D.facilities_text()
    readme = open(os.path.join(target, P.FILES["readme"]), encoding="utf-8").read()
    # SF1: building meters carry month-end reads only, so no day has a metered figure
    with open(os.path.join(target, P.FILES["amr"]), encoding="utf-8") as f:
        rd = [r["read_date"] for r in csv.DictReader(f)]
    ok(all(x.endswith("-01") for x in rd), "SF1 AMR reads are month-end reads")
    ok("counting water lost from the mains as used" in letter, "SF1 the letter counts losses as used")
    # SF2: the largest 24 consecutive hours is the calendar day
    ok(F["rolling"]["figure"] == F["gold_use"], "SF2 rolling 24 hours: %d" % F["rolling"]["figure"])
    ok("(midnight to midnight, Eastern time)" in letter, "SF2 SF3 the letter's calendar day")
    # SF3: UTC days give another figure; the readme says UTC
    ok(F["R2"]["call"] != F["call"], "SF3 UTC days")
    ok("Timestamps are UTC." in readme, "SF3 readme")
    # SF4: a July-to-June year inside the data (October 2025 to June 2026) has no shortfall
    fy = [u for d, u in zip(Bal["days"], Bal["U"]) if d <= dt.date(2026, 6, 30)]
    ok(AN.call_of(max(fy)) is None, "SF4 the fiscal-year window holds no peak")
    ok("twelve months from October 1, 2025 to September 30, 2026" in letter, "SF4 window in the letter")
    # SF5: W1 is the largest under every record
    for d0 in (dt.date(2025, 12, 31), dt.date(2026, 6, 30), dt.date(2026, 10, 12)):
        caps = core.latest_tests(d0)
        ok(max(caps, key=caps.get) == "W1", "SF5 W1 largest as of %s" % d0)
    ok(max(core.NAMEPLATE, key=core.NAMEPLATE.get) == "W1", "SF5 W1 largest nameplate")
    # SF6: firm capacity readings
    at_peak = core.latest_tests(dt.date(2026, 8, 20))
    ok(sum(at_peak.values()) - max(at_peak.values()) == core.FIRM, "SF6 tests in force at the peak converge")
    old = core.latest_tests(dt.date(2025, 9, 30))
    ok(AN.call_of(F["gold_use"] + 0 * 1) is not None and
       int(math.ceil(F["gold_use"] / 1440 - (sum(old.values()) - max(old.values())))) != F["call"],
       "SF6 the 2025 tests give another call")
    nameplate_firm = sum(core.NAMEPLATE.values()) - max(core.NAMEPLATE.values())
    ok(AN.call_of(F["gold_use"] - (nameplate_firm - core.FIRM) * 1440) != F["call"], "SF6 nameplate moves the call")
    ok("each as shown by its most recent pump test" in letter, "SF6 the letter's firm capacity")
    # SF7 and SF8: the gap is not a whole gpm and sits above the half
    ok(0.5 < F["gap"] - math.floor(F["gap"]) < 1.0, "SF7 SF8 gap fraction")
    # SF9: nominal capacity is gallons per foot times height for every tank
    for t in core.TANKS:
        ok(core.nominal(t) == core.GPF[t] * core.HEIGHT[t], "SF9 tank %s" % t)
        ok("{:,}".format(core.GPF[t]) in fac and "{:,}".format(core.nominal(t)) in fac, "SF9 table %s" % t)
    # SF10: every local midnight of the year carries a good sample for each tank
    n, vals, q = P.scada_values(W)
    mids = [i for i, t in enumerate(H["t"][:n]) if core.local(t).hour == 0]
    ok(len(mids) == 366, "SF10 366 local midnights in the export (%d)" % len(mids))
    for tk in core.TANKS:
        ok(all(q["TK%s_LIT" % tk][i] == "G" for i in mids), "SF10 tank %s good at every midnight" % tk)
        for w in core.WELLS:
            pass
    for w in core.WELLS:
        ok(all(q["%s_FQT" % w][i] == "G" for i in mids), "SF10 totalizer %s good at every midnight" % w)
    ok("Levels and flows are instantaneous samples at the timestamp" in readme, "SF10 readme on samples")
    # SF11: the log's totalizers equal the SCADA totalizer tags at every midnight
    wb = load_workbook(os.path.join(target, P.FILES["log"]), read_only=True, data_only=True)
    logged = {}
    for ws in wb.worksheets:
        for r in ws.iter_rows(min_row=6, values_only=True):
            if r and isinstance(r[0], dt.datetime):
                logged[r[0].date()] = r[1:6]
    for i in mids:
        d = core.local(H["t"][i]).date()
        ok(tuple(int(vals["%s_FQT" % w][i]) for w in core.WELLS) == tuple(logged[d]), "SF11 %s" % d)
    ok("the operators record the totalizers at midnight in the daily log" in fac, "SF11 facilities summary")
    # SF12: the daylight-saving days are not the maximum
    for d in (dt.date(2025, 11, 2), dt.date(2026, 3, 8)):
        k = Bal["days"].index(d)
        ok(Bal["U"][k] < F["gold_use"] * 0.9, "SF12 %s" % d)
    # SF13: the system serves the campus alone
    ok("The system serves the campus alone" in fac, "SF13")
    ok(not any(m["bldg"] not in {b["no"] for b in core.B} | {47} for m in A.meters(W)), "SF13 every meter on campus")


def asks_checks(W, final, G, target):
    gold = G["gold"]
    edge = abs((G["share"] * 10) % 1 - 0.5) / 10
    ok(edge >= 0.02, "share %.4f within 0.02 points of its rounding edge" % G["share"])
    nat = G["stops"][(True, True, True, True, "used")]["lost"]
    ok(nat.min() >= 2.5e6, "the sub-meter read leaves every month positive and plausible")
    ok(len(set(gold["lost"])) == 12, "monthly losses all distinct")
    ok(gold["year"] % 1000 != 0 and all(v % 1000 != 0 for v in gold["groups"].values()), "no round totals")
    # every composition leaves every figure it touches
    events = {k for k in range(12) for m in final
              if (m["start"] is not None and A.MONTHS[k] <= m["start"].date() < A.MONTHS[k + 1]) or
              (m["end"] is not None and A.MONTHS[k] <= m["end"].date() < A.MONTHS[k + 1])}
    vault_groups = {A.category(m) for m in final if m["vault"]} - {None}
    event_groups = {A.category(m) for m in final if m["start"] is not None or m["end"] is not None} - {None}
    cells = 0
    for key, st in G["stops"].items():
        subs, vault, partial, mult, use = key
        if key == (False, True, True, True, "used"):
            continue
        touched = set()
        if subs or (not vault) or (not mult):
            touched |= set(range(12)) | {"year", "share"}
        if not partial:
            touched |= events | {"year", "share"}
        if use == "prod":
            touched |= set(range(12)) | {"year"}
        for k in range(12):
            if k in touched:
                ok(int(st["lost"][k]) != gold["lost"][k], "composition %s leaves month %d" % (key, k))
                cells += 1
        if "year" in touched:
            ok(st["year"] != gold["year"], "composition %s leaves the year" % (key,))
            cells += 1
        if "share" in touched:
            ok("%.1f" % st["share"] != G["share_r"], "composition %s leaves the share" % (key,))
            cells += 1
    for key, gr in G["gstops"].items():
        subs, vault, partial, mult, p2 = key
        if key == (False, True, True, True, "golden"):
            continue
        touched = set()
        if subs or not mult:
            touched |= set(A.GROUPS)
        if not vault:
            touched |= vault_groups
        if not partial:
            touched |= event_groups
        if p2 == "current":
            touched |= {"res", "acad"}
        if p2 == "drop":
            touched |= {"acad"}
        for g in touched:
            ok(gr[g] != gold["groups"][g], "composition %s leaves group %s" % (key, g))
            cells += 1
    # the mirror misses every month and the year by the storage change
    prod = G["stops"][(False, True, True, True, "prod")]
    ok(all(abs(int(a) - b) >= 5000 for a, b in zip(prod["lost"], gold["lost"])), "the mirror misses every month")
    # P1 evidence: each sub-metered building has fewer domestic service lines than meters running together
    with open(os.path.join(target, P.FILES["lsl"]), encoding="utf-8") as f:
        lsl = list(csv.DictReader(f))
    lines = {}
    for r in lsl:
        if r["service_use"] not in ("Fire", "Irrigation"):
            lines[int(r["bldg_no"])] = lines.get(int(r["bldg_no"]), 0) + 1
    for parent in {s[1] for s in A.SUBS}:
        running = [m for m in final if m["bldg"] == parent and m["end"] is None and m["start"] is None]
        ok(len(running) > lines[parent], "P1 building %d: %d meters, %d lines" % (parent, len(running), lines[parent]))
    for b in (101, 103, 302, 801):
        ok(lines[b] == 2, "P1 benign twin %d has two service lines" % b)
    ok(sum(1 for m in final if m["sub"]) == 8 and all(m["endpoint"] for m in final if m["sub"]), "P1 subs on AMR")
    # P2 evidence: the annex meter, the removed sheet, Larch Hall's occupancy
    ax = next(m for m in final if m["key"] == "annex")
    ok(ax["bldg"] == 47 and ax["end"] == core.ANNEX_REMOVED, "P2 annex meter on building 47, removed")
    rows = P.meter_rows(final)
    ok(any(r[0] == ax["no"] and "Sci Annex" in r[2] for r in rows), "P2 the annex meter's location")
    ok(sum(1 for m in final if m["bldg"] == 47) == 2, "P2 two meters under building 47")
    ok(G["gstops"][(False, True, True, True, "current")]["res"] - gold["groups"]["res"] ==
       gold["groups"]["acad"] - G["gstops"][(False, True, True, True, "current")]["acad"] > 0, "P2 moves the annex")
    # no building meter's register runs backwards inside the year (no rollover on the ask path)
    for m in final:
        seq = [m["reads"][t] for t in sorted(m["reads"])]
        ok(all(b >= a for a, b in zip(seq, seq[1:])), "register of %s never decreases" % m["no"])
    # H2: vault meters absent from the AMR export and read on work orders, 13 reads each
    amr_meters = {m["no"] for m in final if m["endpoint"]}
    for m in final:
        if m["vault"]:
            ok(m["no"] not in amr_meters and len(m["reads"]) == 13, "H2 vault meter %s" % m["no"])
    # H4: change-out readings agree between the inventory and the work orders
    wo = P.cmms_rows(W, final)
    co = {r[2]: r for r in wo if r[1] == "METER CHANGE-OUT"}
    for k in A.CHANGEOUTS:
        old = next(m for m in final if m["key"] == k and m["service"] and m["line"] == 0 and m["end"] is not None)
        new = next(m for m in final if m["key"] == k and m["service"] and m["line"] == 0 and m["start"] is not None)
        ok(co[new["no"]][7] == old["final_read"] and co[new["no"]][8] == new["initial_read"], "H4 %s" % k)
    return cells


def separation(W, final, target):
    """No device row in the main call's files: the log, the historian export, the tag list, the readme, the pump
    tests, the facilities summary and the letter carry no meter number, sub-meter purpose or building group."""
    main_files = [P.FILES[k] for k in ("log", "tests", "scada", "tags", "readme")]
    meter_nos = {m["no"] for m in final}
    text = D.facilities_text() + "\n" + D.letter_text()
    for f in main_files:
        p = os.path.join(target, f)
        if f.endswith(".xlsx"):
            text += xlsx_strings(p)
        elif f.endswith(".csv") and "scada_historian" in f:
            continue
        else:
            text += open(p, encoding="utf-8").read()
    ok(not any(no in text for no in meter_nos), "separation: no meter number in a main-call file")
    for w in ("make-up", "Concessions", "Pool fill", "Kitchen", "RO/DI", "Vivarium", "vault", "Vault", "AMR",
              "endpoint"):
        ok(w not in text, "separation: %r in a main-call file" % w)
    with open(os.path.join(target, P.FILES["scada"]), encoding="utf-8") as f:
        tags = {line.split(",", 1)[0] for line in f}
    ok(tags == {"tag"} | {t[0] for t in P.TAGS}, "separation: the historian carries the tag list's tags only")


COUNTER = ["storage change", "change in storage", "changes in storage", "adjusted for storage", "net of",
           "drawdown", "drawn down", "draw down", "mass balance", "water balance", "production less", "less storage",
           "refill", "sub-meter", "submeter", "deduct", "downstream", "maximum day demand is", "drew",
           "tank gain", "tanks gained", "gave up"]


def counter_pins(target, final, W):
    texts = {"letter": D.letter_text(), "facilities": D.facilities_text(),
             "readme": open(os.path.join(target, P.FILES["readme"]), encoding="utf-8").read(),
             "guide": open(os.path.join(target, P.FILES["guide"]), encoding="utf-8").read(),
             "tags": open(os.path.join(target, P.FILES["tags"]), encoding="utf-8").read(),
             "log": xlsx_strings(os.path.join(target, P.FILES["log"])),
             "tests": xlsx_strings(os.path.join(target, P.FILES["tests"])),
             "meters": xlsx_strings(os.path.join(target, P.FILES["meters"])),
             "space": xlsx_strings(os.path.join(target, P.FILES["space"])),
             "cmms": "\n".join(" ".join(str(x) for x in r) for r in P.cmms_rows(W, final)),
             "index": open(os.path.join(target, P.FILES["index"]), encoding="utf-8").read()}
    for k, t in texts.items():
        low = t.lower()
        for phrase in COUNTER:
            if phrase == "maximum day demand is" and k == "letter":
                continue
            if phrase == "drawdown" and k == "tests":
                continue          # a pump test's drawdown is the well's, below static water level
            ok(phrase not in low, "counter-pin %r in %s" % (phrase, k))
    ok(texts["letter"].count("Maximum day demand is") == 1, "the definition stated once")
    ok("gallons per foot" not in texts["letter"].lower(), "gallons per foot only in the facilities summary")
    for k, t in texts.items():
        ok("—" not in t and "–" not in t, "no em or en dash in %s" % k)


def realism(W, final, F, G, target, golden):
    gold = G["gold"]
    ok(len(set(gold["used"])) == 12 and len(set(gold["metered"])) == 12, "no identical monthly counts")
    ok(F["gold_use"] % 1000 != 0 and F["R0"]["figure"] % 1000 != 0, "no round headline volumes")
    names = {"Shelly McCarthy", "Wayne Hicks", "Samuel Reyes", "Darrell Clark"}
    docs = D.letter_text() + D.facilities_text()
    found = set(re.findall(r"\b(Shelly McCarthy|Wayne Hicks|Samuel Reyes|Darrell Clark)\b", docs))
    ok(found <= names and "Samuel Reyes" in found and "Shelly McCarthy" in found, "personas from the card")
    # timestamps: nothing in the pack after the extract
    with open(os.path.join(target, P.FILES["cmms"]), encoding="utf-8") as f:
        last = max(dt.datetime.strptime(r["completed"], "%m/%d/%Y %H:%M") for r in csv.DictReader(f))
    ok(last <= dt.datetime(2026, 10, 9, 23, 59), "CMMS within its export")
    with open(os.path.join(target, P.FILES["amr"]), encoding="utf-8") as f:
        ok(max(r["read_date"] for r in csv.DictReader(f)) == "2026-10-01", "AMR within its export")
    ok(len(set(gold["lost"])) == 12, "losses vary month to month")
    # the goldens carry the figures the write-up owes
    from openpyxl import load_workbook as lw
    wb = lw(os.path.join(golden, "campus_water_capacity_2027.xlsx"), data_only=True)
    ev = {r[0]: r[1] for r in wb["Evaluation"].iter_rows(min_row=5, values_only=True) if r and r[0]}
    ok(ev["Water used on the maximum day (gal)"] == F["gold_use"], "golden workbook: water used")
    ok(ev["Well capacity to add before 1 June 2027 (gpm)"] == F["call"], "golden workbook: the call")
    ok(ev["Firm capacity (gpm)"] == core.FIRM, "golden workbook: firm capacity")
    t = docx_text(os.path.join(golden, "sanitary_survey_response_2026.docx"))
    for v in gold["lost"] + [gold["year"]] + list(gold["groups"].values()):
        ok("{:,}".format(v) in t, "golden letter carries %d" % v)
    ok(G["share_r"] + "%" in t and "%d gpm" % F["call"] in t, "golden letter: share and call")


def gates(target, task):
    import json
    files = sorted(os.listdir(target))
    ok(len(files) >= 10, "ten or more files")
    ok(len({f.rsplit(".", 1)[1] for f in files}) >= 3, "three or more formats")
    big = 0
    for f in files:
        if f.endswith(".csv"):
            with open(os.path.join(target, f), encoding="utf-8") as fh:
                big = max(big, sum(1 for _ in fh) - 1)
    ok(big >= 25000, "a file of 25,000 rows or more (%d)" % big)
    meta = json.load(open(os.path.join(task, "metadata.json")))
    ok(len(meta["distractor_files"]) >= 2 and all(d in files for d in meta["distractor_files"]), "distractors")
    for f in files:
        ok(not re.search(r"distractor|golden|answer|solution|trap|decoy", f, re.I), "file name %s" % f)


def run(W, final, F, G, target, golden, task, info):
    N[0] = 0
    ladder(W, F)
    forks(W, F, target)
    cells = asks_checks(W, final, G, target)
    separation(W, final, target)
    counter_pins(target, final, W)
    realism(W, final, F, G, target, golden)
    gates(target, task)
    return N[0], cells
