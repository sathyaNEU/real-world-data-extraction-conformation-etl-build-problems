#!/usr/bin/env python3
"""Independent verifier for task120.

Reads only the shipped bundle (target/): no generator module, no seed, no side file. Every
construction, every calibration outcome and every graded figure is recomputed in DuckDB SQL (the
generator works in pandas) and compared with the expectation table below, which carries the build
record's values. Exit status 0 when every check holds.

    python3 task120/generator/verify_pack.py [path/to/target]
"""
import math
import sys
from pathlib import Path

import duckdb
import openpyxl

EXPECT = {'asks': {'A1_receipts': {'tier1 inst1': 57100430,
                          'tier1 inst2': 60982470,
                          'tier1 inst3': 59690330,
                          'tier1 inst4': 62279690,
                          'tier2 inst1': 27130950,
                          'tier2 inst2': 28344190,
                          'tier2 inst3': 28273480,
                          'tier2 inst4': 29413200,
                          'tier3 inst1': 12597420,
                          'tier3 inst2': 13026520,
                          'tier3 inst3': 13083940,
                          'tier3 inst4': 13625350},
          'A2_net_capital_gain': {'tier1': 3147270635, 'tier2': 995783766, 'tier3': 339052645},
          'A2_share_pct': {'tier1': 30.4, 'tier2': 10.1, 'tier3': 4.5},
          'A3_withholding': {'tier1': 135364253, 'tier2': 229035640, 'tier3': 184720391},
          'tier_agi': {'1': 10356865342, '2': 9837059854, '3': 7492715926}},
 'exact_floors': {"('all', 'ffu', 'attached')": [194429, 288579, 640396],
                  "('all', 'ffu', 'dropped')": [194429, 288579, 600218],
                  "('all', 'ffu', 'separate')": [179741, 268213, 565112],
                  "('all', 'return', 'attached')": [178907, 259351, 573175],
                  "('all', 'return', 'dropped')": [178756, 258646, 559725],
                  "('all', 'return', 'separate')": [166425, 242863, 531325],
                  "('code1', 'ffu', 'attached')": [214063, 318081, 742043],
                  "('code1', 'ffu', 'dropped')": [214063, 318081, 657057],
                  "('code1', 'ffu', 'kept')": [197042, 293068, 639318],
                  "('code1', 'ffu', 'separate')": [197042, 293068, 594090],
                  "('code1', 'return', 'attached')": [195594, 283211, 618397],
                  "('code1', 'return', 'dropped')": [195388, 282266, 592013],
                  "('code1', 'return', 'separate')": [179405, 261583, 551318]},
 'floors_rounded': {"('all', 'ffu', 'attached')": [194000, 289000, 640000],
                    "('all', 'ffu', 'dropped')": [194000, 289000, 600000],
                    "('all', 'ffu', 'separate')": [180000, 268000, 565000],
                    "('all', 'return', 'attached')": [179000, 259000, 573000],
                    "('all', 'return', 'dropped')": [179000, 259000, 560000],
                    "('all', 'return', 'separate')": [166000, 243000, 531000],
                    "('code1', 'ffu', 'attached')": [214000, 318000, 742000],
                    "('code1', 'ffu', 'dropped')": [214000, 318000, 657000],
                    "('code1', 'ffu', 'kept')": [197000, 293000, 639000],
                    "('code1', 'ffu', 'separate')": [197000, 293000, 594000],
                    "('code1', 'return', 'attached')": [196000, 283000, 618000],
                    "('code1', 'return', 'dropped')": [195000, 282000, 592000],
                    "('code1', 'return', 'separate')": [179000, 262000, 551000]},
 'hits': {"('all', 'ffu', 'attached')": 0,
          "('all', 'ffu', 'dropped')": 0,
          "('all', 'ffu', 'separate')": 0,
          "('all', 'return', 'attached')": 0,
          "('all', 'return', 'dropped')": 0,
          "('all', 'return', 'separate')": 0,
          "('code1', 'ffu', 'attached')": 69,
          "('code1', 'ffu', 'dropped')": 55,
          "('code1', 'ffu', 'kept')": 66,
          "('code1', 'ffu', 'separate')": 58,
          "('code1', 'return', 'attached')": 3,
          "('code1', 'return', 'dropped')": 0,
          "('code1', 'return', 'separate')": 3}}  # the build record's values

YEARS = (2022, 2023, 2024, 2025)
CLASSES = [(100_000, 200_000), (200_000, 300_000), (300_000, 500_000), (500_000, 1_000_000),
           (1_000_000, 1_500_000), (1_500_000, 2_000_000), (2_000_000, 5_000_000),
           (5_000_000, 10_000_000), (10_000_000, None)]
CONSTRUCTIONS = [(res, unit, dep) for res in ("code1", "all") for unit in ("ffu", "return")
                 for dep in ("attached", "dropped", "separate")] + [("code1", "ffu", "kept")]


class V:
    def __init__(self):
        self.n = 0
        self.failed = []

    def check(self, cond, name, detail=""):
        self.n += 1
        if not cond:
            self.failed.append(f"{name}: {detail}")


def k_of(n, pct):
    return (n * pct + 99) // 100


def thousand(x):
    return int(math.floor(x / 1000.0 + 0.5)) * 1000


def half_up_k(x):
    q, r = divmod(int(x), 1000)
    return q + (1 if r >= 500 else 0)


def load(con, tgt):
    for y in YEARS:
        con.execute(f"create table r{y} as select * from read_parquet('{tgt}/returns_processed_ty{y}.parquet')")
        con.execute(f"create table s{y} as select * from read_parquet('{tgt}/dependents_schedule_ty{y}.parquet')")
    con.execute(f"create table ledger as select * from read_parquet('{tgt}/estimated_payments_ledger_2025.parquet')")
    con.execute(f"create table returned as select * from read_csv('{tgt}/returned_items_2025.csv', all_varchar=false, header=true)")
    con.execute(f"create table reg as select * from read_csv('{tgt}/tin_match_cases.csv', header=true)")
    con.execute(f"create table schd as select * from read_csv('{tgt}/schedule_d_extract_ty2025.csv', header=true)")
    con.execute(f"create table amended as select * from read_csv('{tgt}/amended_returns_log_ty2025.csv', header=true)")
    con.execute(f"create table efile as select * from read_parquet('{tgt}/wage_statements_efile_ty2025.parquet')")
    rows = []
    for line in open(tgt / "w2_paper_keyed_ty2025.txt", encoding="ascii"):
        rows.append((int(line[10:19]), int(line[19:28]), int(line[43:53]), int(line[32:43])))
    con.execute("create table paper (employer_ein bigint, employee_tin bigint, state_tax_withheld bigint, state_wages bigint)")
    con.executemany("insert into paper values (?, ?, ?, ?)", rows)


def units_sql(y, res, unit, dep):
    rf = "residency_code = 1" if res == "code1" else "true"
    own_key = "r.return_id" if unit == "return" else "r.federal_primary_tin"
    claim_key = "c.return_id" if unit == "return" else "c.federal_primary_tin"
    base = f"""
      with r as (select * from r{y} where {rf}),
      d as (select r.return_id, {claim_key} as ckey
            from r join s{y} s on r.filer_tin = s.dependent_tin join r c on c.return_id = s.claimant_return_id)"""
    if dep == "kept":
        return base + f"""
      , u as (select {own_key} as key, r.federal_agi as agi from r
              union all select d.ckey, r.federal_agi from r join d on r.return_id = d.return_id)
      select key, sum(agi)::bigint as agi from u group by key"""
    if dep == "attached":
        sel = f"coalesce(d.ckey, {own_key})"
        where = "true"
    elif dep == "dropped":
        sel, where = own_key, "d.return_id is null"
    else:
        sel, where = own_key, "true"
    return base + f"""
      select {sel} as key, sum(r.federal_agi)::bigint as agi
      from r left join d on r.return_id = d.return_id where {where} group by 1"""


def county_sql(y, res, unit):
    rf = "residency_code = 1" if res == "code1" else "true"
    if unit == "return":
        return f"select return_id as key, county_code from r{y} where {rf}"
    return f"select filer_tin as key, county_code from r{y} where {rf} and filer_tin = federal_primary_tin"


def floors(con, sql):
    n = con.execute(f"select count(*) from ({sql})").fetchone()[0]
    out = []
    for p in (10, 5, 1):
        k = k_of(n, p)
        vals = [v[0] for v in con.execute(f"select agi from ({sql}) order by agi desc limit 3 offset {k - 2}").fetchall()]
        # vals = ranks k-1, k, k+1; linear interpolation (type 7) at h = (n - 1) * (1 - p/100)
        h = (n - 1) * (1 - p / 100.0)
        frac = h - math.floor(h)
        lin = vals[2] + frac * (vals[1] - vals[2])
        conv = {thousand(vals[0]), thousand(vals[1]), thousand(vals[2]), thousand(lin), thousand((vals[1] + vals[2]) / 2)}
        out.append((vals[1], conv))
    return n, out


def cells(con, y, res, unit, dep, app):
    sql = units_sql(y, res, unit, dep)
    c = {}
    for (lo, hi), i in zip(CLASSES, range(9)):
        cond = f"agi >= {lo}" + (f" and agi < {hi}" if hi else "")
        n, s = con.execute(f"select count(*), coalesce(sum(agi), 0) from ({sql}) where {cond}").fetchone()
        c[("units", i)] = int(n)
        c[("agi_k", i)] = half_up_k(int(s))
    c[("total", 0)] = half_up_k(int(con.execute(f"select sum(agi) from ({sql})").fetchone()[0]))
    if y == 2024:
        csql = county_sql(y, res, unit)
        rows = con.execute(f"select k.county_code, count(*) from ({sql}) u join ({csql}) k on u.key = k.key "
                           f"where u.agi >= 500000 group by 1").fetchall()
        got = dict(rows)
        for code in app:
            c[("county", code)] = int(got.get(code, 0))
    return c


def read_tables(tgt):
    pub, app = {}, []
    for y in (2022, 2023, 2024):
        wb = openpyxl.load_workbook(tgt / f"household_income_tables_ty{y}.xlsx", data_only=True)
        ws = wb["Table 1"]
        c, i = {}, 0
        for row in ws.iter_rows(values_only=True):
            if isinstance(row[0], str) and row[0].startswith("$") and isinstance(row[1], (int, float)):
                c[("units", i)] = int(row[1])
                c[("agi_k", i)] = int(row[2])
                i += 1
            elif isinstance(row[0], str) and row[0].startswith("All full-year resident"):
                c[("total", 0)] = int(row[2])
        if y == 2024:
            for row in wb["Appendix A"].iter_rows(values_only=True):
                if row[1] and str(row[1]).isdigit() and isinstance(row[2], (int, float)):
                    c[("county", str(row[1]))] = int(row[2])
                    app.append(str(row[1]))
        pub[y] = c
    return pub, app


def tiers(con, unit, F):
    """Tier of every TY2025 resident return and the TIN -> tier map, on households (unit 'hh') or
    federal filing units (unit 'ffu')."""
    hk = "coalesce(d.claim_fpt, r.federal_primary_tin)" if unit == "hh" else "r.federal_primary_tin"
    con.execute(f"""create or replace table t_ret as
        with r as (select * from r2025 where residency_code = 1),
        d as (select r.return_id, c.federal_primary_tin as claim_fpt
              from r join s2025 s on r.filer_tin = s.dependent_tin join r c on c.return_id = s.claimant_return_id)
        select r.return_id, r.filer_tin, r.spouse_tin, r.filing_status, r.federal_agi, {hk} as hkey
        from r left join d on r.return_id = d.return_id""")
    con.execute(f"""create or replace table t_hh as select hkey, sum(federal_agi)::bigint as agi,
        case when sum(federal_agi) >= {F[2]} then 1 when sum(federal_agi) >= {F[1]} then 2
             when sum(federal_agi) >= {F[0]} then 3 else 0 end as t from t_ret group by hkey""")
    con.execute("""create or replace table t_member as
        select r.filer_tin as tin, h.t from t_ret r join t_hh h using (hkey)
        union all select r.spouse_tin, h.t from t_ret r join t_hh h using (hkey)
        where r.filing_status = 2 and r.spouse_tin is not null""")
    con.execute("create or replace table t_rid as select r.return_id, h.t from t_ret r join t_hh h using (hkey)")
    return {int(t): int(a) for t, a in con.execute("select t, sum(agi) from t_hh where t > 0 group by t").fetchall()}


def resolved_tier(tin_col, register=True):
    """(expression, joins): the tier of a TIN as reported, through the register's resolved cases."""
    if register:
        return ("coalesce(m.t, m2.t, 0)",
                f"left join t_member m on {tin_col} = m.tin "
                f"left join (select * from reg where case_status = 'RESOLVED') g on {tin_col} = g.reported_tin "
                f"left join t_member m2 on g.resolved_tin = m2.tin")
    return ("coalesce(m.t, 0)", f"left join t_member m on {tin_col} = m.tin")


def receipts(con, clock=True, xfer=True, returned=True, register=True, overclean=False):
    excl = "select txn_id from returned" if overclean else (
        "select txn_id from returned where coalesce(represented_result, '') <> 'PAID'" if returned else "select -1")
    tr = ""
    if not overclean:
        if xfer:
            tr = """union all select l.account_tin, l.amount, cast(o.txn_utc as timestamptz) from ledger l
                    join ledger o on l.source_ref = cast(o.txn_id as varchar)
                    where l.tax_year = 2025 and l.txn_type = 'TRF' and length(l.source_ref) = 11"""
        else:
            tr = """union all select account_tin, amount, cast(txn_utc as timestamptz) from ledger
                    where tax_year = 2025 and txn_type = 'TRF'"""
    if clock:
        when = """case when timezone('America/Chicago', ts) <= timestamp '2025-04-15 23:59:59' then 1
                       when timezone('America/Chicago', ts) <= timestamp '2025-06-16 23:59:59' then 2
                       when timezone('America/Chicago', ts) <= timestamp '2025-09-15 23:59:59' then 3 else 4 end"""
    else:
        when = """case when cast(timezone('UTC', ts) as date) <= date '2025-04-15' then 1
                       when cast(timezone('UTC', ts) as date) <= date '2025-06-15' then 2
                       when cast(timezone('UTC', ts) as date) <= date '2025-09-15' then 3 else 4 end"""
    expr, joins = resolved_tier("p.tin", register)
    sql = f"""
      with p as (select account_tin as tin, amount, cast(txn_utc as timestamptz) as ts from ledger
                 where tax_year = 2025 and txn_type = 'ES' and txn_id not in ({excl}) {tr}),
      q as (select p.amount, p.ts as ts, {expr} as t from p {joins})
      select t, {when} as k, sum(amount)::bigint from q where t > 0 group by all"""
    out = {(t, k): 0 for t in (1, 2, 3) for k in (1, 2, 3, 4)}
    for t, k, s in con.execute(sql).fetchall():
        out[(int(t), int(k))] = int(s)
    return out


def gains(con, versions=True, loss_limit=True, original=False):
    col = "amount_in_agi" if loss_limit else "net_gain_loss"
    if original:
        pick = f"select return_id, {col} as v from amended where amendment_seq = 0 and {col} is not null"
    elif versions:
        pick = (f"select return_id, {col} as v from (select *, row_number() over (partition by return_id order by amendment_seq desc) rn "
                f"from amended where disposition = 'ACCEPTED' and {col} is not null) where rn = 1")
    else:
        pick = "select -1 as return_id, 0 as v"
    sql = f"""with a as ({pick}),
      x as (select e.return_id, coalesce(a.v, e.{col}) as g from schd e left join a on a.return_id = e.return_id)
      select t.t, sum(x.g)::bigint from x join t_rid t using (return_id) where t.t > 0 group by 1"""
    return {int(t): int(s) for t, s in con.execute(sql).fetchall()}


def withholding(con, paper=True, register=True, overclean=False):
    pp = ""
    if paper:
        cond = "where employer_ein not in (select employer_ein from efile)" if overclean else ""
        pp = f"union all select employee_tin, state_tax_withheld from paper {cond}"
    expr, joins = resolved_tier("p.tin", register)
    sql = f"""with p as (select employee_tin as tin, state_tax_withheld as w from efile {pp}),
      q as (select p.w, {expr} as t from p {joins})
      select t, sum(w)::bigint from q where t > 0 group by 1"""
    return {int(t): int(s) for t, s in con.execute(sql).fetchall()}


def main():
    tgt = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / "target").resolve()
    v = V()
    con = duckdb.connect()
    load(con, tgt)
    E = EXPECT
    # ---------------- the published corpus and every construction's reproduction count
    pub, app = read_tables(tgt)
    v.check(sum(len(c) for c in pub.values()) == 69, "corpus.cells")
    hits = {}
    for c in CONSTRUCTIONS:
        h = 0
        for y in (2022, 2023, 2024):
            got = cells(con, y, *c, app)
            h += sum(1 for k, val in pub[y].items() if got.get(k) == val)
        hits[c] = h
        v.check(h == E["hits"][str(c)], f"hits.{c}", (h, E["hits"][str(c)]))
    # the stop rung's misses: the $500K to $1M class and five counties, all short
    ms = []
    for y in (2022, 2023, 2024):
        got = cells(con, y, "code1", "ffu", "separate", app)
        ms += [(y, k, got[k], val) for k, val in pub[y].items() if got[k] != val]
    v.check(len(ms) == 11 and all(g < w and (w - g) / w < 0.05 for _, _, g, w in ms), "stop_rung.misses", ms)
    # ---------------- TY2025 floors under every construction
    for c in CONSTRUCTIONS:
        n, fl = floors(con, units_sql(2025, *c))
        exact = [x[0] for x in fl]
        conv = [x[1] for x in fl]
        v.check(exact == E["exact_floors"][str(c)], f"floors.exact.{c}", (exact, E["exact_floors"][str(c)]))
        v.check(all(len(s) == 1 for s in conv), f"floors.conventions.{c}", conv)
        v.check([thousand(x) for x in exact] == E["floors_rounded"][str(c)], f"floors.rounded.{c}")
    n_hh = con.execute(f"select count(*) from ({units_sql(2025, 'code1', 'ffu', 'attached')})").fetchone()[0]
    n_ffu = con.execute(f"select count(*) from ({units_sql(2025, 'code1', 'ffu', 'separate')})").fetchone()[0]
    v.check((n_hh, n_ffu) == (582_544, 677_351), "units", (n_hh, n_ffu))
    ans = [thousand(x) for x in E["exact_floors"][str(("code1", "ffu", "attached"))]]
    v.check(ans == [214_000, 318_000, 742_000], "answer")
    for F, k in zip(ans, (58_255, 29_128, 5_826)):
        cnt = con.execute(f"select count(*) from ({units_sql(2025, 'code1', 'ffu', 'attached')}) where agi >= {F}").fetchone()[0]
        v.check(cnt == k, f"units_at_floor.{F}", cnt)
    # ---------------- the asks
    tagi = tiers(con, "hh", ans)
    v.check(tagi == {int(k): val for k, val in E["asks"]["tier_agi"].items()}, "tier_agi", tagi)
    g1 = receipts(con)
    for (t, k), val in g1.items():
        v.check(val == E["asks"]["A1_receipts"][f"tier{t} inst{k}"], f"A1.{t}.{k}", val)
    g2 = gains(con)
    for t in (1, 2, 3):
        v.check(g2[t] == E["asks"]["A2_net_capital_gain"][f"tier{t}"], f"A2.{t}", g2[t])
        share = round(100.0 * g2[t] / tagi[t], 1)
        v.check(share == E["asks"]["A2_share_pct"][f"tier{t}"], f"A2share.{t}", share)
        exact = 1000.0 * g2[t] / tagi[t]
        frac = exact % 1
        v.check(abs(frac - 0.5) >= 0.2 and min(frac, 1 - frac) >= 0.1, f"A2share.bin.{t}", exact / 10)
    g3 = withholding(con)
    for t in (1, 2, 3):
        v.check(g3[t] == E["asks"]["A3_withholding"][f"tier{t}"], f"A3.{t}", g3[t])
    # no reported TIN is a listed dependent who does not file, so joining records through the
    # schedule picks up nothing the filer-and-spouse resolution does not (the order of joins cannot fork)
    bad = con.execute("""
        with listed as (select dependent_tin as tin from s2022 union select dependent_tin from s2023
                        union select dependent_tin from s2024 union select dependent_tin from s2025),
             nonfiling as (select tin from listed where tin not in (select filer_tin from r2025)),
             reported as (select employee_tin as tin from efile union all select employee_tin from paper
                          union all select account_tin from ledger union all select reported_tin from reg)
        select count(*) from reported where tin in (select tin from nonfiling)""").fetchone()[0]
    v.check(bad == 0, "reported_tins_clear_of_dependents", bad)
    # every stop off the golden
    stops = {"A1.S1": receipts(con, False, False, False, False), "A1.S2": receipts(con, True, False, False, False),
             "A1.S3": receipts(con, overclean=True)}
    for lab, s in stops.items():
        d = [100.0 * (s[c] - g1[c]) / g1[c] for c in g1]
        v.check(min(abs(x) for x in d) >= 1.0, f"stop.{lab}", d)
    for lab, s in (("A2.S1", gains(con, False, False)), ("A2.S2", gains(con, True, False)), ("A2.S3", gains(con, original=True))):
        v.check(all(abs(s[t] - g2[t]) / g2[t] >= 0.01 for t in (1, 2, 3)), f"stop.{lab}")
    for lab, s in (("A3.S1", withholding(con, False, False)), ("A3.S2", withholding(con, True, False)), ("A3.S3", withholding(con, overclean=True))):
        v.check(all(abs(s[t] - g3[t]) / g3[t] >= 0.01 for t in (1, 2, 3)), f"stop.{lab}")
    fagi = tiers(con, "ffu", [197_000, 293_000, 594_000])
    s4 = (receipts(con), gains(con), withholding(con))
    v.check(all(s4[0][c] != g1[c] for c in g1) and all(s4[1][t] != g2[t] and s4[2][t] != g3[t] for t in (1, 2, 3)), "stop.S4")
    # ---------------- referee, the Department's table, the antidotes
    wb = openpyxl.load_workbook(tgt / "employer_reconciliations_ty2025.xlsx", data_only=True)
    rows = {r[0]: r for r in wb.active.iter_rows(values_only=True) if r[0]}
    ew, pw = con.execute("select (select sum(state_tax_withheld) from efile), (select sum(state_tax_withheld) from paper)").fetchone()
    v.check(rows["Electronic (platform)"][4] == ew and rows["Paper (capture vendor)"][4] == pw, "referee")
    wb = openpyxl.load_workbook(tgt / "returns_processed_by_agi_class_ty2025.xlsx", data_only=True)
    tot = [r for r in wb.active.iter_rows(values_only=True) if r[0] == "All returns"][0]
    n25, a25 = con.execute("select count(*), sum(federal_agi) from r2025").fetchone()
    v.check((tot[1], tot[2]) == (n25, a25), "department_table")
    bad = con.execute("""select count(*) from ledger l join r2024 r on l.source_ref = cast(r.return_id as varchar)
                         where l.txn_type = 'TRF' and l.amount <> r.overpayment_credit_elect""").fetchone()[0]
    v.check(bad == 0, "credit_elect_antidote")
    bad = con.execute("""with a as (select *, row_number() over (partition by return_id order by amendment_seq desc) rn
                         from amended where disposition = 'ACCEPTED')
                         select count(*) from a join r2025 r using (return_id) where a.rn = 1 and a.federal_agi <> r.federal_agi""").fetchone()[0]
    v.check(bad == 0, "version_of_record")
    print(f"verify_pack: {v.n} checks, {len(v.failed)} failed")
    for f in v.failed:
        print("  FAIL", f)
    return 1 if v.failed else 0


if __name__ == "__main__":
    sys.exit(main())
