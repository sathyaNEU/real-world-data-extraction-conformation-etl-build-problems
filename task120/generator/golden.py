"""The two goldens: the source capacity evaluation workbook and the utility's response letter to the district
engineer, with the water audit's bars."""
import datetime as dt
import io
import os

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import asks as A
import core
import docs as D
import postprocess as PP

XLSX = "campus_water_capacity_2027.xlsx"
DOCX = "sanitary_survey_response_2026.docx"
MONTH_NAMES = [m.strftime("%B %Y") for m in A.MONTHS[:12]]
GROUP_NAMES = {"res": "Residence halls", "acad": "Academic and research buildings", "ath": "Athletics facilities",
               "cep": "Central energy plant"}
BOLD = Font(name="Calibri", bold=True)
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(name="Calibri", bold=True, color="FFFFFF")
KEY_FILL = PatternFill("solid", fgColor="FCE9C8")
THIN = Side(style="thin", color="B7C3D0")


def g(x):
    return "{:,}".format(int(x))


def workbook(path, F, Gf, W):
    wb = Workbook()
    ws = wb.active
    ws.title = "Evaluation"
    ws["A1"] = "%s (PWSID %s)" % (core.SYSTEM, core.PWSID)
    ws["A1"].font = Font(name="Calibri", bold=True, size=13)
    ws["A2"] = "Source capacity evaluation for the 2026 sanitary survey: 1 October 2025 to 30 September 2026"
    ws["A2"].font = Font(name="Calibri", italic=True)
    rows = [
        ("Maximum day (Eastern time, midnight to midnight)", F["gold_day"], "mmmm d, yyyy", "Daily use sheet"),
        ("Water used on the maximum day (gal)", F["gold_use"], "#,##0", "Production less storage change"),
        ("    Well production that day (gal)", F["gold_P"], "#,##0", "Daily log, meter totalizers at 00:00"),
        ("    Water the four tanks gave up that day (gal)", -F["gold_dS"], "#,##0",
         "SCADA levels at 00:00 x gallons per foot"),
        ("Maximum day demand (gpm)", F["gold_use"] / 1440.0, "#,##0.0", "Gallons a day / 1,440"),
        ("Firm capacity (gpm)", F["firm"], "#,##0", "Wells sheet"),
        ("Well capacity to add before 1 June 2027 (gpm)", F["call"], "#,##0", "Rounded up to a whole gpm"),
    ]
    ws.append([])
    ws.append(["Measure", "Figure", "Source"])
    for c in ws[4]:
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
    for label, v, fmt, srcl in rows:
        ws.append([label, v, srcl])
        ws.cell(row=ws.max_row, column=3).font = Font(name="Calibri", italic=True, color="555555")
        c = ws.cell(row=ws.max_row, column=2)
        c.number_format = fmt
        c.alignment = Alignment(horizontal="right")
        if label.startswith("Well capacity to add") or label.startswith("Water used on"):
            for cc in ws[ws.max_row]:
                cc.font = BOLD
                cc.fill = KEY_FILL
    ws.append([])
    ws.append(["Firm capacity is the total of the wells' most recent pump-test capacities less Well 1, the "
               "largest: " + ", ".join("%s %s" % (w, g(core.TESTED[w])) for w in core.WELLS) + " gpm."])
    ws.append(["Maximum day demand is the day's well production less the water the four tanks gained between that "
               "day's midnights (here they gave water up), counting water lost from the mains as used."])
    ws.append(["The largest day of well production alone was %s (%s gal), when the wells refilled the tanks."
               % (F["R0"]["day"].strftime("%B %-d"), g(F["R0"]["figure"]))])
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 40
    ws.freeze_panes = "A5"

    ws2 = wb.create_sheet("Daily use")
    ws2.append(["Date", "Well production (gal)", "Storage at 00:00 (gal)", "Storage change (gal)",
                "Water used (gal)"])
    for c in ws2[1]:
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    Bal = W["Bal"]
    S0 = Bal["S_all"]
    for k, d in enumerate(Bal["days"]):
        ws2.append([d, int(Bal["P"][k]), int(S0[Bal["by_date"][d]]), int(Bal["dS"][k]), int(Bal["U"][k])])
        if d == F["gold_day"]:
            for cc in ws2[ws2.max_row]:
                cc.fill = KEY_FILL
                cc.font = BOLD
    ws2.append(["Year", int(Bal["P"].sum()), None, int(Bal["dS"].sum()), int(Bal["U"].sum())])
    for cc in ws2[ws2.max_row]:
        cc.font = BOLD
        cc.border = Border(top=THIN)
    for row in ws2.iter_rows(min_row=2):
        row[0].number_format = "mm/dd/yyyy"
        for c in row[1:]:
            c.number_format = "#,##0;-#,##0"
    for j, w in enumerate([12, 18, 18, 16, 16], start=1):
        ws2.column_dimensions[get_column_letter(j)].width = w
    ws2.freeze_panes = "B2"
    ws2.auto_filter.ref = "A1:E%d" % (len(Bal["days"]) + 1)
    ws2.print_title_rows = "1:1"

    ws3 = wb.create_sheet("Wells")
    ws3.append(["Well", "Most recent pump test", "Tested capacity (gpm)", "In firm capacity"])
    for c in ws3[1]:
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
    last = {}
    for w, d, gpm in core.PUMP_TESTS:
        if w not in last or d > last[w][0]:
            last[w] = (d, gpm)
    for w in core.WELLS:
        ws3.append([core.WELL_NAMES[w], last[w][0], last[w][1], "No, largest well" if w == core.LARGEST else "Yes"])
    ws3.append(["Firm capacity", None, core.FIRM, None])
    for cc in ws3[ws3.max_row]:
        cc.font = BOLD
    for row in ws3.iter_rows(min_row=2):
        row[1].number_format = "mm/dd/yyyy"
        row[2].number_format = "#,##0"
    for j, w in enumerate([28, 20, 20, 18], start=1):
        ws3.column_dimensions[get_column_letter(j)].width = w

    ws4 = wb.create_sheet("Notes")
    for line in [
        "Evaluation period: 1 October 2025 to 30 September 2026; days run midnight to midnight, Eastern time.",
        "Well production: each well's discharge meter totalizer at the day's two midnights (daily log); the Well 3 "
        "register rolled over on 4 March 2026 and is counted through.",
        "Storage: each tank's level at the day's two midnights (SCADA historian, hourly samples, UTC) times its "
        "gallons per foot of level (facilities summary), summed over Tanks A to D.",
        "Water used on a day: production less the change in storage; it includes the water lost from the mains.",
        "Capacities in gpm: gallons a day divided by 1,440. The capacity to add is the gap rounded up to a whole gpm.",
        "Prepared by Campus Utilities, October 2026."]:
        ws4.append([line])
    ws4.column_dimensions["A"].width = 120
    wb.save(path)
    PP.normalize_xlsx(path, "Shelly McCarthy", "2026-10-13T13:05:00Z", "2026-10-16T09:48:00Z")


def chart_png(Gf):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    lost = Gf["gold"]["lost"]
    labels = [m.strftime("%b\n%Y") if m.month in (10, 1) else m.strftime("%b") for m in A.MONTHS[:12]]
    fig, ax = plt.subplots(figsize=(7.2, 3.6), dpi=150)
    bars = ax.bar(range(12), [x / 1e6 for x in lost], color="#2E6F9E", width=0.68)
    for b, v in zip(bars, lost):
        ax.text(b.get_x() + b.get_width() / 2, b.get_height() + 0.18, "{:,}".format(v), ha="center", va="bottom",
                fontsize=6.3, rotation=90)
    ax.set_xticks(range(12))
    ax.set_xticklabels(labels, fontsize=8)
    ax.set_ylabel("Water lost (million gallons)", fontsize=8.5)
    ax.set_ylim(0, max(lost) / 1e6 * 1.32)
    ax.tick_params(axis="y", labelsize=8)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.set_title("Water lost by month, October 2025 to September 2026: %s gallons, %s%% of the water used"
                 % ("{:,}".format(Gf["gold"]["year"]), Gf["share_r"]), fontsize=8.8, loc="left")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def _footer(d, text):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    p = d.sections[0].footer.paragraphs[0]
    r = p.add_run(text)
    r.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    run.append(t)
    fld.append(run)
    p._p.append(fld)


def letter(path, F, Gf):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
    gold = Gf["gold"]
    d = Document()
    stl = d.styles["Normal"]
    stl.font.name = "Calibri"
    stl.font.size = Pt(10.5)
    _footer(d, "PWSID 4149327, 2026 Sanitary Survey response, page ")
    p = d.add_paragraph()
    r = p.add_run("Kettle Ridge University")
    r.bold = True
    r.font.size = Pt(13)
    d.add_paragraph("Facilities and Operations, Campus Utilities\n1200 Fraser Street, Kettle Ridge, PA 16899")
    d.add_paragraph("October 16, 2026")
    d.add_paragraph("Samuel Reyes, P.E.\nSanitary Engineer\nPennsylvania Department of Environmental Protection\n"
                    "Safe Drinking Water Program, Northcentral District Office")
    p = d.add_paragraph()
    p.add_run("Re: 2026 Sanitary Survey response, Kettle Ridge University Campus Water System, PWSID 4149327").bold = True
    d.add_paragraph("Dear Mr. Reyes:")
    d.add_paragraph("Thank you for your letter of September 15. Our responses to the three requests follow; the source "
                    "capacity evaluation workbook (campus_water_capacity_2027.xlsx) is attached.")
    d.add_heading("1. Source capacity", level=2)
    d.add_paragraph(
        "The system's maximum day demand in the evaluation period was %s gallons, used on %s (%s gpm). That day the "
        "wells produced %s gallons, with Well 1 out of service for rehabilitation, and the four tanks gave up %s "
        "gallons during the first-year move-in and the heat. Firm capacity, from the wells' most recent pump tests "
        "less Well 1, the largest, is %s gpm, %s gpm short of the maximum day demand."
        % (g(F["gold_use"]), F["gold_day"].strftime("%B %-d, %Y"), "{:,.1f}".format(F["gold_use"] / 1440.0),
           g(F["gold_P"]), g(-F["gold_dS"]), g(F["firm"]), "{:,.1f}".format(F["gap"])))
    p = d.add_paragraph()
    p.add_run("The University will add %d gpm of firm well capacity before June 1, 2027." % F["call"]).bold = True
    p.add_run(" Capital Planning has scheduled the replacement of the Well 2 pump and motor with a larger unit: design "
              "in November 2026, installation in April 2027 and a pump test by May 15, 2027, which we will send to the "
              "district office when it is complete. Well 1 remains the largest well.")
    d.add_heading("2. Water audit", level=2)
    d.add_paragraph("Water lost is the water used in the system less the water delivered to the campus's buildings as "
                    "their meters recorded it. For the year the system lost %s gallons of the %s gallons used, a "
                    "water loss percentage of %s%%." % (g(gold["year"]), g(gold["used_year"]), Gf["share_r"]))
    tb = d.add_table(rows=1 + 12 + 1, cols=4)
    tb.style = "Light Grid Accent 1"
    hdr = ["Month", "Water used (gal)", "Delivered to buildings (gal)", "Water lost (gal)"]
    for j, h in enumerate(hdr):
        tb.cell(0, j).text = h
    for k in range(12):
        vals = [MONTH_NAMES[k], g(gold["used"][k]), g(gold["metered"][k]), g(gold["lost"][k])]
        for j, v in enumerate(vals):
            tb.cell(k + 1, j).text = v
    tot = ["Year", g(gold["used_year"]), g(sum(gold["metered"])), g(gold["year"])]
    for j, v in enumerate(tot):
        tb.cell(13, j).text = v
    for i, row in enumerate(tb.rows):
        for j, cell in enumerate(row.cells):
            for pp in cell.paragraphs:
                if j > 0:
                    pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for run in pp.runs:
                    run.font.size = Pt(9)
                    run.bold = i in (0, 13)
    src_p = d.add_paragraph("Water used: the wells' meters less the change in the four tanks' storage between the "
                            "month's first midnights. Delivered: the buildings' service meters, AMR month-end reads "
                            "and the vault meters' logged reads.")
    src_p.runs[0].italic = True
    src_p.runs[0].font.size = Pt(8.5)
    d.add_picture(io.BytesIO(chart_png(Gf)), width=Inches(6.3))
    d.add_paragraph("Grounds irrigation is unmetered and counts in the water lost; so does the water used to flush "
                    "hydrants in April and lost at the two main breaks in January and February.")
    d.add_heading("3. Water use by building group", level=2)
    d.add_paragraph("Water delivered over the evaluation period, as the buildings' meters recorded it:")
    tb2 = d.add_table(rows=5, cols=2)
    tb2.style = "Light Grid Accent 1"
    tb2.cell(0, 0).text = "Building group"
    tb2.cell(0, 1).text = "Water used (gal)"
    for i, k in enumerate(A.GROUPS, start=1):
        tb2.cell(i, 0).text = GROUP_NAMES[k]
        tb2.cell(i, 1).text = g(gold["groups"][k])
    for i, row in enumerate(tb2.rows):
        for j, cell in enumerate(row.cells):
            for pp in cell.paragraphs:
                if j == 1:
                    pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for run in pp.runs:
                    run.font.size = Pt(9.5)
                    run.bold = i == 0
    d.add_paragraph("The academic and research buildings include the Science Annex until its meter was removed on "
                    "December 19, 2025, before the building was demolished; the residence halls include Larch Hall, "
                    "which took over the annex's building number, from its meter's installation on July 14, 2026.")
    d.add_paragraph("Please call me with any questions.")
    d.add_paragraph("Sincerely,\n\nShelly McCarthy\nDirector of Campus Utilities")
    d.add_paragraph("cc: Darrell Clark, Capital Planning; Wayne Hicks, Chief Water Operator")
    cp = d.core_properties
    cp.author = "Shelly McCarthy"
    cp.last_modified_by = "Shelly McCarthy"
    cp.title = "2026 Sanitary Survey response, PWSID 4149327"
    cp.comments = ""
    cp.keywords = ""
    cp.subject = ""
    cp.category = ""
    cp.revision = 3
    cp.created = dt.datetime(2026, 10, 14, 8, 20)
    cp.modified = dt.datetime(2026, 10, 16, 10, 2)
    buf = io.BytesIO()
    d.save(buf)
    D.normalize_docx(buf.getvalue(), path)


def write(golden, F, Gf, W):
    os.makedirs(golden, exist_ok=True)
    workbook(os.path.join(golden, XLSX), F, Gf, W)
    letter(os.path.join(golden, DOCX), F, Gf)
