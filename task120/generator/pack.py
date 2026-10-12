"""Write the data files of target/: the SCADA historian export and its tag list, the operators' daily log, the
pump tests, the AMR month-end reads, the meter inventory, the CMMS work orders, the space inventory, the 2024
service line inventory, and the two neighbouring files (the weather station's daily record and the grounds
irrigation programs)."""
import csv
import datetime as dt
import os

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import asks as A
import core
import postprocess as PP

FILES = {
    "scada": "scada_historian_export_2025-10-01_2026-10-01.csv",
    "tags": "scada_tag_list.csv",
    "readme": "scada_export_readme.txt",
    "log": "water_operators_daily_log_2025-10_2026-10.xlsx",
    "tests": "well_pump_tests_2021-2026.xlsx",
    "facilities": "campus_water_system_facilities_summary_2026.docx",
    "letter": "dep_sanitary_survey_letter_2026-09-15.pdf",
    "amr": "amr_monthend_reads_2025-10_2026-10.csv",
    "meters": "meter_inventory_2026-10-08.xlsx",
    "cmms": "cmms_water_work_orders_2025-10-01_2026-10-09.csv",
    "space": "space_inventory_buildings_2026-09-30.xlsx",
    "lsl": "service_line_inventory_2024-10-15.csv",
    "weather": "weather_station_daily_2025-10_2026-09.csv",
    "irrigation": "grounds_irrigation_programs_2026.xlsx",
    "ccr": "krbwa_2025_drinking_water_quality_report.pdf",
    "index": "survey_response_records_index.csv",
    "guide": "campus_water_records_field_guide.txt",
}
DISTRACTORS = ["weather", "irrigation", "ccr"]
EXPORT_END = dt.datetime(2026, 10, 2, 4, 0, tzinfo=core.UTC)     # exclusive
TAGS = []
for w in core.WELLS:
    n = w[1]
    TAGS += [("%s_FIT" % w, "Well %s discharge flow" % n, "gpm", "Well field", "instantaneous"),
             ("%s_FQT" % w, "Well %s discharge meter totalizer" % n, "gal", "Well field", "cumulative"),
             ("%s_RUN" % w, "Well %s pump running" % n, "1/0", "Well field", "status")]
TAGS += [("TKA_LIT", "Standpipe No. 1 level", "ft", "Main zone", "instantaneous"),
         ("TKB_LIT", "Stadium Elevated Tank level", "ft", "Main zone", "instantaneous"),
         ("TKC_LIT", "North Campus Reservoir level", "ft", "Hill zone", "instantaneous"),
         ("TKD_LIT", "North Campus Elevated Tank level", "ft", "Hill zone", "instantaneous"),
         ("BPS_P1_RUN", "North Campus Booster Station pump 1 running", "1/0", "Booster station", "status"),
         ("BPS_P2_RUN", "North Campus Booster Station pump 2 running", "1/0", "Booster station", "status"),
         ("BPS_FIT", "Booster station discharge flow", "gpm", "Booster station", "instantaneous"),
         ("BPS_PIT_S", "Booster station suction pressure", "psi", "Booster station", "instantaneous"),
         ("BPS_PIT_D", "Booster station discharge pressure", "psi", "Booster station", "instantaneous"),
         ("MZ_PIT", "Main zone pressure, Old Main", "psi", "Main zone", "instantaneous"),
         ("HZ_PIT", "Hill zone pressure, Health Center", "psi", "Hill zone", "instantaneous"),
         ("EP_CL2", "Entry point free chlorine residual", "mg/L", "Well field", "instantaneous")]
HILL_RTU = {"TKC_LIT", "TKD_LIT", "BPS_P1_RUN", "BPS_P2_RUN", "BPS_FIT", "BPS_PIT_S", "BPS_PIT_D", "HZ_PIT"}
RTU_OUT = [dt.datetime(2026, 3, 17, h, 0, tzinfo=core.UTC) for h in (18, 19, 20)]
TKD_FAULT = [dt.datetime(2026, 6, 2, 15, 0, tzinfo=core.UTC)]
HEAD = Font(name="Calibri", bold=True)
FILL = PatternFill("solid", fgColor="DDE6F0")


def _w(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(header)
        w.writerows(rows)


def fmt_num(v, nd):
    return ("%%.%df" % nd) % v


def scada_values(W):
    H = W["H"]
    n = sum(1 for t in H["t"] if t < EXPORT_END)
    out = {}
    for j, w in enumerate(core.WELLS):
        out["%s_FIT" % w] = [fmt_num(x, 1) for x in H["flow"][w][:n]]
        out["%s_FQT" % w] = ["%d" % x for x in H["tot"][w][:n]]
        out["%s_RUN" % w] = ["%d" % x for x in H["run"][:n, j]]
    for tk in core.TANKS:
        out["TK%s_LIT" % tk] = [fmt_num(x, 2) for x in H["lev"][tk][:n]]
    out["BPS_P1_RUN"] = ["%d" % x for x in H["prun"][:n, 0]]
    out["BPS_P2_RUN"] = ["%d" % x for x in H["prun"][:n, 1]]
    out["BPS_FIT"] = [fmt_num(x, 1) for x in H["bflow"][:n]]
    out["BPS_PIT_S"] = [fmt_num(x, 1) for x in H["suct"][:n]]
    out["BPS_PIT_D"] = [fmt_num(x, 1) for x in H["disch"][:n]]
    out["MZ_PIT"] = [fmt_num(x, 1) for x in H["mz"][:n]]
    out["HZ_PIT"] = [fmt_num(x, 1) for x in H["hz"][:n]]
    out["EP_CL2"] = [fmt_num(x, 2) for x in H["cl"][:n]]
    q = {k: ["G"] * n for k in out}
    pos = {t: i for i, t in enumerate(H["t"][:n])}
    for t in RTU_OUT:
        i = pos[t]
        for k in HILL_RTU:
            out[k][i] = out[k][pos[RTU_OUT[0]] - 1]
            q[k][i] = "U"
    for t in TKD_FAULT:
        out["TKD_LIT"][pos[t]] = "0.00"
        q["TKD_LIT"][pos[t]] = "B"
    rng = np.random.default_rng(2718)
    groups = {"well field": [t for t in out if t.startswith("W") or t == "EP_CL2"],
              "booster": [t for t in out if t.startswith("BPS")], "zone pressure": ["MZ_PIT", "HZ_PIT"]}
    times = H["t"][:n]
    for k in range(14):
        g = list(groups)[k % 3]
        while True:
            i = int(rng.integers(30, n - 30))
            span = int(rng.integers(1, 4))
            hrs = [times[i + s] for s in range(span + 1)]
            if all(x.hour != 0 and core.local(x).hour != 0 for x in hrs):
                break
        for s in range(1, span + 1):
            for tag in groups[g]:
                out[tag][i + s] = out[tag][i]
                q[tag][i + s] = "U"
    for tag in ("W2_FIT", "W4_FIT", "BPS_PIT_D", "MZ_PIT", "EP_CL2", "HZ_PIT"):
        while True:
            i = int(rng.integers(30, n - 30))
            if times[i].hour != 0 and core.local(times[i]).hour != 0 and q[tag][i] == "G":
                break
        out[tag][i] = "0.0" if tag != "EP_CL2" else "0.00"
        q[tag][i] = "B"
    return n, out, q


def outage_days(W):
    """The local dates of the well-field radio losses in the historian (for the operators' remarks)."""
    n, vals, q = scada_values(W)
    H = W["H"]
    days = sorted({core.local(H["t"][i]).date() for i in range(n) if q["W1_FQT"][i] == "U"})
    return days[::2]


def write_scada(W, target):
    H = W["H"]
    n, vals, q = scada_values(W)
    rows = []
    for i in range(n):
        ts = H["t"][i].strftime("%Y-%m-%dT%H:%M:%SZ")
        for tag, *_ in TAGS:
            rows.append((tag, ts, vals[tag][i], q[tag][i]))
    _w(os.path.join(target, FILES["scada"]), ["tag", "timestamp_utc", "value", "quality"], rows)
    _w(os.path.join(target, FILES["tags"]), ["tag", "description", "units", "area", "value_type", "note"],
       [(t, d, u, a, v, "8-digit register: rolls over after 99,999,999" if t == "W3_FQT" else "")
        for t, d, u, a, v in TAGS])
    return n


def _sheet(ws, header, rows, widths=None, freeze="A2", fmt=None):
    ws.append(header)
    for c in ws[1]:
        c.font = HEAD
        c.fill = FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(list(r))
    if freeze:
        ws.freeze_panes = freeze
    for j, w in enumerate(widths or [14] * len(header), start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    if fmt:
        for col, f in fmt.items():
            for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                for c in row:
                    c.number_format = f


LOG_REMARKS = {
    dt.date(2025, 11, 2): "Clocks back 02:00",
    dt.date(2026, 1, 22): "6-in main break College Ave at Fraser St; isolated 14:20, repaired 22:40",
    dt.date(2026, 2, 11): "W4 out of service all day, check valve replaced",
    dt.date(2026, 2, 26): "4-in main break Reservoir Rd 06:30, repaired 11:50",
    dt.date(2026, 3, 4): "W3 totalizer rolled over (8-digit register)",
    dt.date(2026, 3, 8): "Clocks forward 02:00",
    dt.date(2026, 3, 17): "Hill RTU radio down approx 14:00 to 17:00, values held",
    dt.date(2026, 6, 2): "TKD level transmitter fault 11:00, reset by I&C",
    dt.date(2026, 8, 10): "W1 out of service for rehabilitation, contractor mobilized",
    dt.date(2026, 8, 12): "BPS pump 1 out, seal replacement",
    dt.date(2026, 8, 26): "BPS pump 1 back in service",
    dt.date(2026, 8, 29): "W1 back in service after rehabilitation",
    dt.date(2026, 9, 15): "W1 pump test after rehab, 1,090 gpm (contractor report to follow)",
}
for d in core.FLUSHING:
    LOG_REMARKS[d] = "Unidirectional flushing, main zone, 08:00 to 15:00"
LOG_REMARKS[dt.date(2026, 8, 26)] = "DEP sanitary survey on site (S. Reyes), 09:00 to 14:30"
ROUTINE = ["EP Cl2 analyzer calibrated", "Hypochlorite delivered, W2 house", "Hypochlorite delivered, W3 house",
           "Hypochlorite delivered, W1 house", "Booster station PM, pumps greased", "W2 pump PM, oil checked",
           "W4 house heater checked", "Heavy rain overnight", "Chlorine residual low at Hillside Hall tap, flushed H-41",
           "Valve crew on College Ave", "Lab courier late, Bac-T samples held in cooler",
           "Annual pressure gauge calibration, all well houses"]
ONE_OFF = {dt.date(2026, 2, 2): "Annual cross-connection survey letters mailed",
           dt.date(2025, 10, 20): "Leak reported near Old Main lawn, WO opened",
           dt.date(2026, 5, 18): "Leak reported near Old Main lawn again, WO opened",
           dt.date(2026, 1, 6): "Lead and copper sampling kits delivered to 20 sites"}
HOLIDAYS = {dt.date(2025, 11, 27), dt.date(2025, 11, 28), dt.date(2025, 12, 24), dt.date(2025, 12, 25),
            dt.date(2026, 1, 1), dt.date(2026, 1, 19), dt.date(2026, 5, 25), dt.date(2026, 7, 3), dt.date(2026, 9, 7)}


def w5_exercise_days():
    """Well 5 is exercised to waste on the first Tuesday of each month (the meter sits after the waste tee)."""
    out, d = [], core.EVAL0
    while d < core.EXTRACT:
        if d.weekday() == 1 and d.day <= 7:
            out.append(d)
        d += dt.timedelta(1)
    return out


def remarks(W, outages=()):
    rng = np.random.default_rng(808)
    out = dict(LOG_REMARKS)
    out.update(ONE_OFF)
    for d0 in outages:
        out.setdefault(d0, "Well field RTU comms lost for a few hours, values held")
    H = W["H"]
    mids = sorted((core.local(t).date(), i) for i, t in enumerate(H["t"]) if core.local(t).hour == 0)
    for (d0, i0), (d1, i1) in zip(mids[:-1], mids[1:]):
        if H["tot"]["W3"][i1] < H["tot"]["W3"][i0]:
            out[d0] = "W3 totalizer rolled over (8-digit register)"
    d = core.EVAL0
    while d < core.EXTRACT:
        if d.weekday() == 0 and d not in HOLIDAYS and d not in out:
            out[d] = "Bac-T samples collected, %s routes" % ("north" if (d.isocalendar()[1] % 2) else "south")
        elif d not in out and d.weekday() < 5 and d not in HOLIDAYS and rng.random() < 0.09:
            out[d] = ROUTINE[int(rng.integers(0, len(ROUTINE)))]
        d += dt.timedelta(1)
    return out


def w5_exercise_minutes():
    rng = np.random.default_rng(55)
    return {d: int(rng.choice([15, 20, 20, 25, 30])) for d in w5_exercise_days()}


def write_log(W, target):
    H, S, G = W["H"], W["S"], W["G"]
    rng = np.random.default_rng(1203)
    loc_mid = {core.local(t).date(): i for i, t in enumerate(H["t"]) if core.local(t).hour == 0}
    at8 = {core.local(t).date(): i for i, t in enumerate(H["t"]) if core.local(t).hour == 8}
    runh = {}
    for k, d in enumerate(G["days"]):
        m = G["di"] == k
        runh[d] = (S["run"][m].sum(0) * 0.25, S["prun"][m].sum(0) * 0.25)
    rem = remarks(W, outage_days(W))
    w5x = w5_exercise_minutes()
    for d, mins in w5x.items():
        rem[d] = ((rem[d] + "; ") if d in rem else "") + "W5 exercised to waste, %d min" % mins
    wb = Workbook()
    wb.remove(wb.active)
    d = core.EVAL0
    months = []
    while d < core.EXTRACT:
        if (d.year, d.month) not in months:
            months.append((d.year, d.month))
        d += dt.timedelta(1)
    for y, m in months:
        ws = wb.create_sheet(dt.date(y, m, 1).strftime("%b %Y"))
        ws.append(["%s, PWSID %s" % (core.SYSTEM, core.PWSID)])
        ws.append(["Operators' daily log, %s" % dt.date(y, m, 1).strftime("%B %Y")])
        ws.append(["Meter readings are the well discharge meter totalizers at 00:00 on the date. Run hours are "
                   "for the date (00:00 to 24:00)."])
        ws.append([])
        head = (["Date"] + ["%s meter (gal)" % w for w in core.WELLS] + ["%s hrs" % w for w in core.WELLS]
                + ["BPS P1 hrs", "BPS P2 hrs", "Cl2 EP (mg/L)", "Remarks"])
        ws.append(head)
        for c in ws[5]:
            c.font = HEAD
            c.fill = FILL
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws["A1"].font = Font(name="Calibri", bold=True, size=12)
        dd = dt.date(y, m, 1)
        while dd.month == m and dd < core.EXTRACT:
            i = loc_mid[dd]
            wh, bh = runh[dd]
            wh = [float(x) for x in wh]
            if dd in w5x:
                wh[4] += round(w5x[dd] / 60.0, 1)
            cl = float(np.round(W["H"]["cl"][at8[dd]] + rng.normal(0, 0.03), 2))
            row = ([dd] + [int(H["tot"][w][i]) for w in core.WELLS] + [round(x, 1) for x in wh]
                   + [round(float(x), 1) for x in bh] + [cl, rem.get(dd, None)])
            ws.append(row)
            dd += dt.timedelta(1)
        ws.freeze_panes = "B6"
        widths = [11] + [14] * 5 + [7] * 5 + [9, 9, 9, 52]
        for j, wd in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(j)].width = wd
        for row in ws.iter_rows(min_row=6):
            row[0].number_format = "mm/dd/yyyy"
            for c in row[1:6]:
                c.number_format = "#,##0"
            for c in row[6:13]:
                c.number_format = "0.0"
            row[13].number_format = "0.00"
    path = os.path.join(target, FILES["log"])
    wb.save(path)
    PP.normalize_xlsx(path, "Wayne Hicks", "2025-10-01T07:12:00Z", "2026-10-12T11:40:00Z")


def write_tests(target):
    rng = np.random.default_rng(77)
    wb = Workbook()
    ws = wb.active
    ws.title = "Pump tests"
    rows = []
    static = {"W1": 62.0, "W2": 71.0, "W3": 58.0, "W4": 83.0, "W5": 66.0}
    for w, d, g in core.PUMP_TESTS:
        sw = static[w] + rng.normal(0, 2.2) + (4.0 if d.year == 2026 else 0)
        sc = g / (38.0 + rng.normal(0, 2.5)) if not (w == "W1" and d == dt.date(2026, 4, 21)) else g / 52.0
        if w == "W1" and d == dt.date(2026, 9, 15):
            sc = g / 31.0
        dd = g / sc
        note = ""
        if w == "W1" and d == dt.date(2026, 4, 21):
            note = "Capacity down from 2025; rehabilitation scheduled for August 2026"
        if w == "W1" and d == dt.date(2026, 9, 15):
            note = "After rehabilitation (brushing, airlift, acid treatment; new bowl assembly)"
        rows.append((core.WELL_NAMES[w], w, d, "Centre Pump & Well Co.", g, 24 if d.year >= 2023 else 8,
                     round(sw, 1), round(sw + dd, 1), round(dd, 1), round(sc, 1), note))
    rows.sort(key=lambda r: (r[1], r[2]))
    _sheet(ws, ["Well", "Well ID", "Test date", "Tested by", "Pumping rate (gpm)", "Duration (h)",
                "Static water level (ft bgs)", "Pumping water level (ft bgs)", "Drawdown (ft)",
                "Specific capacity (gpm/ft)", "Notes"], rows,
           widths=[26, 8, 11, 22, 11, 9, 12, 12, 10, 12, 60], fmt={3: "mm/dd/yyyy", 5: "#,##0"})
    ws2 = wb.create_sheet("Wells")
    drilled = {"W1": 1962, "W2": 1967, "W3": 1974, "W4": 1988, "W5": 1996}
    depth = {"W1": 412, "W2": 388, "W3": 455, "W4": 501, "W5": 436}
    hp = {"W1": 125, "W2": 100, "W3": 100, "W4": 100, "W5": 75}
    _sheet(ws2, ["Well ID", "Well", "Year drilled", "Depth (ft)", "Pump", "Motor (hp)", "Nameplate (gpm)"],
           [(w, core.WELL_NAMES[w], drilled[w], depth[w], "Vertical turbine", hp[w], core.NAMEPLATE[w])
            for w in core.WELLS], widths=[8, 26, 12, 10, 18, 10, 14])
    path = os.path.join(target, FILES["tests"])
    wb.save(path)
    PP.normalize_xlsx(path, "Wayne Hicks", "2021-05-20T13:02:00Z", "2026-09-21T15:31:00Z")


LOCATIONS = ["Basement mech room", "Ground floor mech room", "Lower level mech room", "Basement, NE corner",
             "Basement, SW corner", "Basement, NW corner", "Basement, SE corner", "Boiler room", "Pump room",
             "Janitor closet, ground floor", "Crawl space access", "Loading dock utility room"]
VAULT_LOC = ["Meter pit, east lawn", "Vault, Fraser St sidewalk", "Meter pit, north service drive",
             "Vault, loading dock", "Meter pit, west lawn", "Vault, College Ave side", "Meter pit, south lot",
             "Vault, service court", "Meter pit, Fraser St side"]
PIT_AMR = {205, 210, 226, 405, 603, 802}       # meters in pits that the network can hear


def meter_rows(final):
    rng = np.random.default_rng(311)
    rows = []
    vi = 0
    pair_loc = {}
    for m in sorted(final, key=lambda x: x["no"]):
        if m["sub"]:
            loc = m["location"]
        elif m["key"] == "annex":
            loc = "Sci Annex, bsmt mech"
        elif m["vault"]:
            loc = VAULT_LOC[vi % len(VAULT_LOC)]
            vi += 1
        elif m["bldg"] in PIT_AMR and m["line"] == 0:
            loc = rng.choice(["Meter pit, front lawn", "Meter pit, east side", "Vault, north side"])
        elif m["line"] == 1:
            loc = rng.choice(LOCATIONS) + ", second service"
        else:
            loc = rng.choice(LOCATIONS + ["Mech room %s%02d" % (rng.choice(["B", "G", "0", "1"]),
                                                                int(rng.integers(1, 40)))] * 4)
        if m["service"] and m["key"] in A.CHANGEOUTS and m["line"] == 0:
            loc = pair_loc.setdefault(m["key"], loc)
        model = A.model_of(rng, m["size"], m["installed"].year)
        status = "Removed" if m["end"] is not None else "Active"
        removed = m["end"].date() if m["end"] is not None else None
        rows.append([m["no"], m["bldg"], loc, m["size"], model, m["mult"], m["purpose"], status, m["installed"],
                     m["initial_read"], removed, m["final_read"]])
    return rows


def write_meters(final, target):
    wb = Workbook()
    ws = wb.active
    ws.title = "Meters"
    rows = meter_rows(final)
    _sheet(ws, ["Meter No.", "Bldg No.", "Location", "Size (in)", "Make/model", "Register multiplier", "Purpose",
                "Status", "Installed", "Initial read", "Removed", "Final read"], rows,
           widths=[10, 9, 32, 8, 24, 10, 24, 9, 11, 11, 11, 13],
           fmt={4: "0.0#", 9: "mm/dd/yyyy", 10: "#,##0", 11: "mm/dd/yyyy", 12: "#,##0"})
    ws.auto_filter.ref = "A1:L%d" % (len(rows) + 1)
    ws2 = wb.create_sheet("Notes")
    for line in ["Campus Utilities meter shop inventory, as of 8 October 2026.",
                 "One row per meter in service on or after 1 October 2025 (meters removed earlier are archived).",
                 "Register multiplier: register reading x multiplier = gallons.",
                 "Initial read: register reading when the meter was set. Final read: register reading when it was "
                 "removed.",
                 "Bldg No. is the space inventory building number the meter was set up under."]:
        ws2.append([line])
    ws2.column_dimensions["A"].width = 110
    path = os.path.join(target, FILES["meters"])
    wb.save(path)
    PP.normalize_xlsx(path, "Campus Utilities Meter Shop", "2019-03-11T14:20:00Z", "2026-10-08T12:05:00Z")


def amr_rows(final):
    rows = []
    for m in final:
        if m["endpoint"] is None:
            continue
        for t, r in m["reads"].items():
            rows.append((m["endpoint"], m["no"], t.date().isoformat(), "00:00", r))
    rows.sort(key=lambda r: (r[0], r[2], r[1]))
    return rows, 0


def write_amr(final, target):
    rows, dup = amr_rows(final)
    _w(os.path.join(target, FILES["amr"]), ["endpoint_id", "meter_no", "read_date", "read_time", "register_read"],
       rows)
    return len(rows), dup


def space_rows():
    rows = []
    area = {"hill": "North Campus"}
    for b in sorted(core.B, key=lambda x: x["no"]):
        a = area.get(b["zone"], "Central Campus" if b["no"] < 400 or b["no"] > 600 else "South Campus")
        note = "Occupied 08/17/2026" if b["no"] == 47 else ""
        rows.append([b["no"], b["name"], b["cat"], a, b["gsf"], b["built"], note])
    for k in range(1, 6):
        rows.append([510 + k, "Well House %d" % k, "Utilities", "South Campus" if k in (4, 5) else "Central Campus",
                     900 + 40 * k, [1962, 1967, 1974, 1988, 1996][k - 1], ""])
    rows.append([516, "North Campus Booster Station", "Utilities", "North Campus", 2400, 1984, ""])
    rows.sort(key=lambda r: r[0])
    return rows


def write_space(target):
    wb = Workbook()
    ws = wb.active
    ws.title = "Buildings"
    _sheet(ws, ["Bldg No.", "Building name", "Primary use", "Campus area", "GSF", "Year built", "Notes"], space_rows(),
           widths=[9, 34, 16, 15, 10, 10, 22], fmt={5: "#,##0"})
    ws2 = wb.create_sheet("Removed from inventory")
    a = core.ANNEX
    _sheet(ws2, ["Bldg No.", "Building name", "Primary use", "GSF", "Year built", "Removed", "Reason"],
           [[a["no"], a["name"], a["cat"], a["gsf"], a["built"], core.ANNEX_REMOVED.date(),
             "Demolished for the Larch Hall project"],
            [318, "Poultry Research Barn", "Research", 6400, 1949, dt.date(2024, 6, 28), "Demolished"],
            [609, "Fraser Street House", "Administrative", 3100, 1923, dt.date(2023, 8, 15), "Sold"]],
           widths=[9, 26, 14, 9, 10, 11, 40], fmt={4: "#,##0", 6: "mm/dd/yyyy"})
    path = os.path.join(target, FILES["space"])
    wb.save(path)
    PP.normalize_xlsx(path, "Office of Space Management", "2016-01-19T15:00:00Z", "2026-09-30T19:22:00Z")


def lsl_rows(final):
    """The 2024 lead service line inventory: each building's service lines as of 15 October 2024 (the annex in,
    Larch Hall not yet built)."""
    rng = np.random.default_rng(902)
    rows = []
    k = 0
    blds = [b for b in core.B if b["no"] != 47] + [core.ANNEX]
    size = {}
    for m in final:
        if m["service"] and m["start"] is None:
            size[(m["key"], m["line"])] = m["size"]
    for b in sorted(blds, key=lambda x: x["no"]):
        key = "annex" if b is core.ANNEX else b["no"]
        lines = []
        for j in range(b["lines"]):
            use = "Process" if (b["no"] == 302 and j == 1) else "Domestic"
            if b["no"] == 501:
                use = "Domestic and process"
            lines.append((use, size[(key, j)]))
        if b["gsf"] >= 60000 or b["no"] in (501, 401, 402, 301):
            lines.append(("Fire", 8.0 if b["gsf"] >= 150000 or b["no"] == 501 else 6.0))
        if b["no"] in (601, 701, 202, 215, 703):
            lines.append(("Irrigation", 1.5))
        for use, sz in lines:
            k += 1
            yr = int(rng.integers(max(b["built"], 1950), max(b["built"], 1950) + 6)) if b["built"] < 2019 else b["built"]
            mat = "Ductile iron" if sz >= 4 else ("Copper" if yr >= 1960 else rng.choice(["Copper", "Galvanized"]))
            cls = "Non-lead" if mat != "Galvanized" else "Galvanized requiring replacement"
            rows.append(["SL-%04d" % (1000 + 7 * k), b["no"], b["name"], use, sz, mat,
                         mat if rng.random() < 0.8 else "Copper", cls, yr,
                         rng.choice(["Records review", "Records review", "Visual inspection at meter",
                                     "Excavation"])])
    return rows


def write_lsl(final, target):
    _w(os.path.join(target, FILES["lsl"]),
       ["service_line_id", "bldg_no", "bldg_name", "service_use", "size_in", "system_side_material",
        "customer_side_material", "classification", "install_year", "basis"], lsl_rows(final))


def cmms_rows(W, final):
    rng = np.random.default_rng(5150)
    rows = []

    def business_day_after(d, k):
        x = d
        while k > 0:
            x += dt.timedelta(1)
            if x.weekday() < 5:
                k -= 1
        return x

    def workday(d0, span):
        while True:
            d = d0 + dt.timedelta(int(rng.integers(0, span)))
            if d.weekday() < 5 and d not in HOLIDAYS:
                return d

    for m in final:
        if not m["vault"]:
            continue
        for t, r in sorted(m["reads"].items()):
            visit = business_day_after(t.date() - dt.timedelta(1), int(rng.integers(1, 4)))
            vt = dt.datetime(visit.year, visit.month, visit.day, int(rng.integers(7, 11)), int(rng.integers(0, 60)))
            rows.append(dict(type="METER READ", asset=m["no"], bldg=m["bldg"],
                             desc="Monthly read, vault meter (handheld download of logged register)",
                             opened=vt.date(), completed=vt, r1=r, r2=None, rfor=t, crew="PLB-2"))
    olds = {m["key"]: m for m in final if m["service"] and m["key"] in A.CHANGEOUTS and m["end"] is not None}
    news = {m["key"]: m for m in final if m["service"] and m["key"] in A.CHANGEOUTS and m["start"] is not None}
    for k, tc in A.CHANGEOUTS.items():
        o, nw = olds[k], news[k]
        op = tc.date() - dt.timedelta(9)
        while op.weekday() >= 5 or op in HOLIDAYS:
            op -= dt.timedelta(1)
        rows.append(dict(type="METER CHANGE-OUT", asset=nw["no"], bldg=k,
                         desc="Replace meter %s with %s (%s-in, scheduled replacement)" % (o["no"], nw["no"],
                                                                                         ("%g" % nw["size"])),
                         opened=op, completed=tc, r1=o["final_read"], r2=nw["initial_read"],
                         rfor=None, crew="PLB-1"))
    ax = next(m for m in final if m["key"] == "annex")
    rows.append(dict(type="METER REMOVAL", asset=ax["no"], bldg=47,
                     desc="Remove meter and cut service, Science Annex (building demolition)",
                     opened=dt.date(2025, 12, 8), completed=core.ANNEX_REMOVED, r1=ax["final_read"], r2=None,
                     rfor=None, crew="PLB-1"))
    lh = next(m for m in final if m["key"] == 47)
    rows.append(dict(type="METER INSTALL", asset=lh["no"], bldg=47,
                     desc="Set 4-in meter and AMR endpoint, Larch Hall (new residence hall)",
                     opened=dt.date(2026, 7, 1), completed=core.LARCH_INSTALL, r1=None, r2=lh["initial_read"],
                     rfor=None, crew="PLB-1"))
    # the rest of the year's water work
    blds = [b for b in core.B if b["no"] != 47]
    for b in blds:
        if b["gsf"] >= 40000:
            while True:
                d = dt.date(2025, 10, 1) + dt.timedelta(int(rng.integers(0, 360)))
                if d.weekday() < 5 and d not in HOLIDAYS:
                    break
            rows.append(dict(type="BACKFLOW TEST", asset="BFP-%d" % b["no"], bldg=b["no"],
                             desc="Annual RPZ test, domestic service", opened=d,
                             completed=dt.datetime(d.year, d.month, d.day, int(rng.integers(7, 15)),
                                                   int(rng.integers(0, 60))),
                             r1=None, r2=None, rfor=None, crew="PLB-3"))
    for m in final:
        if m["service"] and m["size"] >= 3.0 and m["end"] is None and m["start"] is None and rng.random() < 0.35:
            d = workday(dt.date(2025, 10, 6), 340)
            rows.append(dict(type="METER TEST", asset=m["no"], bldg=m["bldg"],
                             desc="Large meter accuracy test, within AWWA limits, returned to service", opened=d,
                             completed=dt.datetime(d.year, d.month, d.day, int(rng.integers(8, 15)),
                                                   int(rng.integers(0, 60))), r1=None, r2=None, rfor=None,
                             crew="PLB-1"))
    for a, b, gpm, z, txt in core.MAIN_BREAKS:
        rows.append(dict(type="MAIN BREAK", asset="MAIN", bldg=None, desc=txt + ", repaired", opened=a.date(),
                         completed=b, r1=None, r2=None, rfor=None, crew="WTR-1"))
    for d in core.FLUSHING:
        rows.append(dict(type="HYDRANT FLUSHING", asset="HYD", bldg=None,
                         desc="Unidirectional flushing, main zone sequence %d" % (core.FLUSHING.index(d) + 1),
                         opened=d, completed=dt.datetime(d.year, d.month, d.day, int(rng.integers(13, 17)),
                                                         int(rng.integers(0, 60))), r1=None, r2=None, rfor=None,
                         crew="WTR-2"))
    for k in range(46):
        d = workday(dt.date(2025, 10, 2), 362)
        rows.append(dict(type="VALVE EXERCISE", asset="V-%03d" % int(rng.integers(1, 260)), bldg=None,
                         desc=["Valve exercised, %d turns, left open", "Valve exercised, %d turns, stiff, left open",
                               "Valve found partly closed, exercised %d turns, left open",
                               "Valve box cleaned, valve exercised %d turns"][int(rng.integers(0, 4))]
                         % int(rng.integers(12, 34)), opened=d,
                         completed=dt.datetime(d.year, d.month, d.day, int(rng.integers(8, 15)),
                                               int(rng.integers(0, 60))), r1=None,
                         r2=None, rfor=None, crew="WTR-2"))
    for k in range(7):
        b = blds[int(rng.integers(0, len(blds)))]
        d = dt.date(2025, 10, 9) + dt.timedelta(int(rng.integers(0, 350)))
        rows.append(dict(type="SERVICE LEAK", asset="SVC-%d" % b["no"], bldg=b["no"],
                         desc=["Service line leak at curb stop, repaired",
                               "Service line leak at the meter pit, coupling replaced",
                               "Service line leak under sidewalk, clamp installed"][int(rng.integers(0, 3))],
                         opened=d,
                         completed=dt.datetime(d.year, d.month, d.day, int(rng.integers(10, 18)),
                                               int(rng.integers(0, 60))) + dt.timedelta(days=int(rng.integers(0, 2))),
                         r1=None, r2=None, rfor=None, crew="WTR-1"))
    fixed = [("WELL REHAB", "W1", None, "Well 1 rehabilitation (contractor): pull pump, brush, airlift, acid, "
              "new bowl assembly", dt.date(2026, 8, 10), dt.datetime(2026, 8, 28, 16, 30)),
             ("PUMP REPAIR", "W4", None, "Well 4 check valve replaced", dt.date(2026, 2, 11),
              dt.datetime(2026, 2, 11, 16, 0)),
             ("PUMP REPAIR", "BPS-P1", 516, "Booster pump 1 mechanical seal replaced", dt.date(2026, 8, 12),
              dt.datetime(2026, 8, 25, 15, 10)),
             ("TANK INSPECTION", "TK-C", None, "North Campus Reservoir ROV inspection, in service", dt.date(2026, 6, 9),
              dt.datetime(2026, 6, 9, 14, 0)),
             ("INSTRUMENT", "TKD_LIT", None, "North Campus Elevated Tank level transmitter fault, reset",
              dt.date(2026, 6, 2), dt.datetime(2026, 6, 2, 12, 5)),
             ("INSTRUMENT", "RTU-HILL", 516, "Hill RTU radio, replaced coax connector", dt.date(2026, 3, 17),
              dt.datetime(2026, 3, 17, 17, 20))]
    for tp, asset, bldg, desc, op, comp in fixed:
        rows.append(dict(type=tp, asset=asset, bldg=bldg, desc=desc, opened=op, completed=comp, r1=None, r2=None,
                         rfor=None, crew="WTR-1"))
    # endpoint batteries, check reads on high-use alerts, hydrant meters, calls and generator tests
    amr_meters = [m for m in final if m["endpoint"] and m["end"] is None and m["start"] is None]
    for k in range(14):
        m = amr_meters[int(rng.integers(0, len(amr_meters)))]
        d = workday(dt.date(2025, 10, 14), 330)
        rows.append(dict(type="ENDPOINT BATTERY", asset=m["no"], bldg=m["bldg"],
                         desc="AMR endpoint low-battery alarm, battery pack replaced", opened=d,
                         completed=dt.datetime(d.year, d.month, d.day, int(rng.integers(8, 15)),
                                               int(rng.integers(0, 60))), r1=None,
                         r2=None, rfor=None, crew="PLB-2"))
    for k in range(8):
        m = amr_meters[int(rng.integers(0, len(amr_meters)))]
        while True:
            d = dt.date(2025, 10, 9) + dt.timedelta(int(rng.integers(0, 335)))
            if 3 <= d.day <= 27 and d.weekday() < 5:
                break
        t = dt.datetime(d.year, d.month, d.day, int(rng.integers(8, 15)), 15 * int(rng.integers(0, 4)))
        rows.append(dict(type="CHECK READ", asset=m["no"], bldg=m["bldg"],
                         desc="Check read on a high-use alert, no leak found", opened=d, completed=t,
                         r1=m["reg"](A.step_index(t)), r2=None, rfor=t, crew="PLB-2"))
    for m in final:
        if m["sub"] and m["location"].startswith("CT deck") and m["endpoint"]:
            t = dt.datetime(2025, 12, 3, 9, 15 + 15 * ("CT-2" in m["location"]))
            rows.append(dict(type="CHECK READ", asset=m["no"], bldg=m["bldg"],
                             desc="Check read on a zero-use alert, towers drained for the winter, no fault",
                             opened=t.date(), completed=t, r1=m["reg"](A.step_index(t)), r2=None, rfor=t,
                             crew="PLB-2"))
    for d in (dt.date(2025, 10, 20), dt.date(2026, 5, 18)):
        rows.append(dict(type="SERVICE LEAK", asset="SVC-601", bldg=601,
                         desc="Leak on Old Main lawn at the 2-in irrigation tap, repaired", opened=d,
                         completed=dt.datetime(d.year, d.month, d.day, 15, 40) + dt.timedelta(days=1), r1=None,
                         r2=None, rfor=None, crew="WTR-1"))
    calls = ["Discolored water call, flushed nearest hydrant", "Low pressure call, top floor, PRV checked",
             "Leak reported at curb box, checked, no leak", "Taste and odor call, chlorine residual checked",
             "Water off for renovation tie-in, notice posted"]
    for k in range(11):
        b = blds[int(rng.integers(0, len(blds)))]
        while True:
            d = dt.date(2025, 10, 3) + dt.timedelta(int(rng.integers(0, 360)))
            if not (dt.date(2026, 8, 17) <= d <= dt.date(2026, 8, 24)):
                break
        rows.append(dict(type="CUSTOMER REQUEST", asset="BLDG-%d" % b["no"], bldg=b["no"],
                         desc=calls[int(rng.integers(0, len(calls)))], opened=d,
                         completed=dt.datetime(d.year, d.month, d.day, int(rng.integers(9, 17)),
                                               int(rng.integers(0, 60))), r1=None,
                         r2=None, rfor=None, crew="WTR-2"))
    for k in range(10):
        w = core.WELLS[k % 5]
        d = dt.date(2025, 10, 7) + dt.timedelta(int(36 * k + rng.integers(0, 20)))
        while d.weekday() >= 5 or d in HOLIDAYS:
            d += dt.timedelta(1)
        rows.append(dict(type="GENERATOR TEST", asset="GEN-%s" % w, bldg=511 + core.WELLS.index(w),
                         desc="Standby generator run under load, 30 min", opened=d,
                         completed=dt.datetime(d.year, d.month, d.day, int(rng.integers(9, 14)),
                                               int(rng.integers(0, 60))), r1=None, r2=None, rfor=None,
                         crew="ELEC"))
    rows.sort(key=lambda r: (r["opened"], r["completed"], r["type"], str(r["asset"])))
    out = []
    for i, r in enumerate(rows):
        out.append(["26-%04d" % (1180 + 3 * i + int(rng.integers(0, 3))) if r["opened"] >= dt.date(2026, 1, 1)
                    else "25-%04d" % (3410 + 3 * i + int(rng.integers(0, 3))),
                    r["type"], r["asset"], "" if r["bldg"] is None else r["bldg"], r["desc"],
                    r["opened"].strftime("%m/%d/%Y"), r["completed"].strftime("%m/%d/%Y %H:%M"),
                    "" if r["r1"] is None else r["r1"], "" if r["r2"] is None else r["r2"],
                    "" if r["rfor"] is None else r["rfor"].strftime("%m/%d/%Y %H:%M"), r["crew"], "CLOSED"])
    return out


def write_cmms(W, final, target):
    rows = cmms_rows(W, final)
    _w(os.path.join(target, FILES["cmms"]),
       ["wo_no", "wo_type", "asset_id", "bldg_no", "description", "opened", "completed", "reading_1", "reading_2",
        "reading_for", "crew", "status"], rows)
    return len(rows)


def write_weather(W, target):
    rng = np.random.default_rng(66)
    rows = []
    for d in core.eval_days():
        k = W["G"]["dix"][d]
        tx = float(W["T"][k])
        tn = tx - 17.5 + rng.normal(0, 3.5)
        p = 0.0 if not W["rain"][k] else float(np.round(rng.gamma(1.1, 0.28), 2))
        rows.append((d.isoformat(), "%.1f" % tx, "%.1f" % tn, "%.2f" % p,
                     "%d" % int(np.clip(68 + rng.normal(0, 9) + (14 if p > 0 else 0), 25, 100)),
                     "%.1f" % max(0.4, rng.normal(6.2, 2.4))))
    _w(os.path.join(target, FILES["weather"]), ["date", "tmax_f", "tmin_f", "precip_in", "rh_mean_pct",
                                                "wind_mean_mph"], rows)


def write_irrigation(target):
    rng = np.random.default_rng(4040)
    wb = Workbook()
    ws = wb.active
    ws.title = "Programs 2026"
    rows = []
    ctrls = [("C1", "Old Main lawns"), ("C2", "Central quad"), ("C3", "Library mall"), ("C4", "Student Union plaza"),
             ("C5", "South Campus fields"), ("C6", "North Campus quad"), ("C7", "College Ave frontage"),
             ("C8", "Arboretum beds")]
    for c, area in ctrls:
        for z in range(1, int(rng.integers(6, 11))):
            kind = rng.choice(["Rotor", "Spray", "Drip"], p=[0.55, 0.35, 0.10])
            gpm = {"Rotor": rng.uniform(22, 48), "Spray": rng.uniform(9, 24), "Drip": rng.uniform(2, 7)}[kind]
            mins = {"Rotor": rng.integers(18, 36), "Spray": rng.integers(8, 16), "Drip": rng.integers(30, 61)}[kind]
            rows.append([c, area, z, kind, round(float(gpm), 1), int(mins), "A" if kind != "Drip" else "B",
                         "22:00" if kind != "Drip" else "04:30", "M W F" if kind == "Rotor" else "Daily",
                         int(rng.choice([80, 90, 100, 110]))])
    _sheet(ws, ["Controller", "Area", "Zone", "Head type", "Zone flow (gpm)", "Run time (min)", "Program",
                "Start", "Days", "Seasonal adjust (%)"], rows, widths=[10, 24, 6, 10, 12, 12, 9, 8, 10, 14],
           fmt={5: "0.0"})
    ws2 = wb.create_sheet("Season")
    for line in ["Grounds Department irrigation programs, 2026 season.",
                 "Controllers run from about 20 April to 10 October; rain sensors suspend a night's programs.",
                 "Zone flows are design flows from the 2019 audit by the irrigation contractor."]:
        ws2.append([line])
    ws2.column_dimensions["A"].width = 100
    path = os.path.join(target, FILES["irrigation"])
    wb.save(path)
    PP.normalize_xlsx(path, "Grounds Department", "2019-04-02T12:44:00Z", "2026-04-17T18:10:00Z")
