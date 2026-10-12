"""The documents of target/: the state's sanitary survey letter, the system's facilities summary, the historian
export's readme, the records field guide and index, and the borough authority's 2025 water quality report."""
import csv
import datetime as dt
import io
import re
import zipfile

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import core
import postprocess as PP

LETTER_TEXT = [
    ("p", "September 15, 2026"),
    ("p", "Shelly McCarthy<br/>Director of Campus Utilities<br/>Kettle Ridge University, Facilities and Operations"
          "<br/>1200 Fraser Street<br/>Kettle Ridge, PA 16899"),
    ("b", "Re: 2026 Sanitary Survey<br/>Kettle Ridge University Campus Water System, PWSID 4149327<br/>"
          "Centre County"),
    ("p", "Dear Ms. McCarthy:"),
    ("p", "The Department completed a sanitary survey of the Kettle Ridge University Campus Water System on August 26, "
          "2026. Well 1 was out of service for rehabilitation at the time of the survey. The other facilities were "
          "inspected and no significant deficiencies were found. To complete the survey, please submit the following "
          "by November 30, 2026."),
    ("h", "1. Source capacity"),
    ("p", "Submit an evaluation of the system's firm capacity against its maximum day demand for the evaluation "
          "period. The evaluation period is the twelve months from October 1, 2025 to September 30, 2026."),
    ("p", "Maximum day demand is the most water used in the system in any one calendar day (midnight to midnight, "
          "Eastern time) of the evaluation period, counting water lost from the mains as used."),
    ("p", "Firm capacity is the total of the wells' capacities, each as shown by its most recent pump test, less the "
          "capacity of the largest well."),
    ("p", "Capacities are compared in gallons per minute, a day's water in gallons divided by 1,440. If firm capacity "
          "is below maximum day demand, state the source capacity, in gallons per minute, that the system will add "
          "before June 1, 2027 so that its firm capacity meets the maximum day demand, with the schedule for the "
          "work."),
    ("h", "2. Water audit"),
    ("p", "Report the water lost in each month of the evaluation period and over the year, and the year's water loss "
          "percentage."),
    ("p", "Water lost in a month is the water used in the system in the month less the water delivered to the "
          "campus's buildings in the month, as their meters recorded it. The water loss percentage is the year's "
          "water lost as a percentage of the water used in the system in the year."),
    ("h", "3. Water use by building group"),
    ("p", "Report the water used over the evaluation period by the campus's residence halls, its academic and "
          "research buildings, its athletics facilities and its central energy plant."),
    ("p", "The system's next sanitary survey is due in 2029. If you have questions about this letter, please contact "
          "me at the district office."),
    ("p", "Sincerely,"),
    ("p", "Samuel Reyes, P.E.<br/>Sanitary Engineer<br/>Safe Drinking Water Program, Northcentral District"),
    ("s", "cc: Wayne Hicks, Chief Water Operator, Kettle Ridge University<br/>Darrell Clark, Capital Planning, "
          "Kettle Ridge University"),
]

FACILITIES = [
    ("t", "Campus Water System: Facilities Summary"),
    ("s", "Kettle Ridge University, Facilities and Operations, Campus Utilities. Revised July 2026."),
    ("h", "The system"),
    ("p", "The Kettle Ridge University Campus Water System (PWSID 4149327) is a community water system owned and "
          "operated by the University. The system serves the campus alone: the residence halls, the academic, "
          "research and athletics buildings, the central energy plant and the University's other buildings. It has "
          "no connection to the Kettle Ridge Borough Water Authority's system. The campus is supplied by five wells "
          "and served in two pressure zones."),
    ("h", "Sources"),
    ("TABLE_WELLS", ""),
    ("p", "Each well's discharge meter totalizes the water it pumps; the operators record the totalizers at midnight "
          "in the daily log. Sodium hypochlorite is fed at each well house, and the combined entry point residual is "
          "monitored continuously. Pump tests are run each spring by the University's well contractor; the results "
          "are kept in the pump test records."),
    ("h", "Storage tanks"),
    ("TABLE_TANKS", ""),
    ("p", "Level is the depth of water above level zero: the floor of the standpipe and the reservoir, the bottom of "
          "the bowl of each elevated tank. The bowls of the two elevated tanks are cylindrical above level zero."),
    ("h", "Pressure zones and the booster station"),
    ("p", "The wells pump into the main zone, which floats on Standpipe No. 1 and the Stadium Elevated Tank "
          "(overflow 1,160 ft). The hill zone, north campus above about 1,090 ft, floats on the North Campus "
          "Reservoir and the North Campus Elevated Tank (overflow 1,262 ft) and is fed from the main zone by the "
          "North Campus Booster Station: two 905-gpm pumps, one duty and one standby, alternating monthly."),
    ("h", "Monitoring and records"),
    ("p", "The water SCADA system polls the well houses, the four tanks, the booster station and two zone pressure "
          "points; the historian keeps an hourly sample of every tag. The tag list gives each tag's units. The "
          "operators keep a daily log of the meter totalizers, pump run hours, the entry point chlorine residual and "
          "remarks."),
    ("p", "Director of Campus Utilities: Shelly McCarthy. Chief Water Operator: Wayne Hicks."),
]


def _styles():
    ss = getSampleStyleSheet()
    return dict(
        body=ParagraphStyle("b", parent=ss["Normal"], fontName="Times-Roman", fontSize=11, leading=14.5, spaceAfter=9),
        bold=ParagraphStyle("bb", parent=ss["Normal"], fontName="Times-Bold", fontSize=11, leading=14.5, spaceAfter=9),
        head=ParagraphStyle("h", parent=ss["Normal"], fontName="Times-Bold", fontSize=11, leading=14.5, spaceAfter=4,
                            spaceBefore=4),
        title=ParagraphStyle("t", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=14, leading=18,
                             spaceAfter=4),
        small=ParagraphStyle("s", parent=ss["Normal"], fontName="Times-Roman", fontSize=9.5, leading=12,
                             spaceAfter=10),
        lh=ParagraphStyle("lh", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=10, leading=12.5,
                          textColor=colors.HexColor("#1F3A5F")),
        lh2=ParagraphStyle("lh2", parent=ss["Normal"], fontName="Helvetica", fontSize=8.5, leading=11,
                           textColor=colors.HexColor("#1F3A5F"), spaceAfter=16),
        cell=ParagraphStyle("c", parent=ss["Normal"], fontName="Helvetica", fontSize=8.4, leading=10.2),
    )


def _footer(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Times-Roman", 8)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(1.0 * inch, 0.6 * inch, text)
        canvas.drawRightString(LETTER[0] - 1.0 * inch, 0.6 * inch, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def letter_pdf(path):
    st = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=LETTER, leftMargin=1.0 * inch, rightMargin=1.0 * inch,
                            topMargin=0.8 * inch, bottomMargin=0.9 * inch, title="2026 Sanitary Survey, PWSID 4149327",
                            author="Safe Drinking Water Program", creator="Safe Drinking Water Program",
                            producer="Safe Drinking Water Program")
    story = [Paragraph("Pennsylvania Department of Environmental Protection", st["lh"]),
             Paragraph("Safe Drinking Water Program, Northcentral District Office", st["lh2"])]
    for kind, text in LETTER_TEXT:
        sty = {"b": st["bold"], "h": st["head"], "s": st["small"]}.get(kind, st["body"])
        story.append(Paragraph(text, sty))
    doc.build(story, onFirstPage=_footer("Safe Drinking Water Program, Northcentral District Office"),
              onLaterPages=_footer("PWSID 4149327, 2026 Sanitary Survey"))
    open(path, "wb").write(buf.getvalue())
    PP.normalize_pdf(path, b"D:20260915103400-04'00'", "Safe Drinking Water Program")


def letter_text():
    return "\n".join(re.sub(r"<br/>", "\n", t) for _, t in LETTER_TEXT)


def _wells_table():
    rows = [["Well", "Name", "Year drilled", "Depth (ft)", "Nameplate (gpm)"]]
    drilled = {"W1": 1962, "W2": 1967, "W3": 1974, "W4": 1988, "W5": 1996}
    depth = {"W1": 412, "W2": 388, "W3": 455, "W4": 501, "W5": 436}
    for w in core.WELLS:
        rows.append([w, core.WELL_NAMES[w].split(" (")[1].rstrip(")"), str(drilled[w]), str(depth[w]),
                     "{:,}".format(core.NAMEPLATE[w])])
    return rows


def _tanks_table():
    rows = [["Tank", "Name", "Zone", "Type", "Built", "Level zero (ft)", "Overflow (ft)", "Height to overflow (ft)",
             "Gallons per foot of level", "Nominal capacity (gal)"]]
    kind = {"A": "Steel standpipe", "B": "Composite elevated", "C": "Prestressed concrete ground",
            "D": "Steel elevated, fluted column"}
    built = {"A": 1961, "B": 1998, "C": 1984, "D": 2009}
    for t in core.TANKS:
        rows.append([t, core.TANK_NAMES[t], "Main" if core.ZONE_OF_TANK[t] == "main" else "Hill", kind[t],
                     str(built[t]), "{:,.0f}".format(core.BASE[t]), "{:,.0f}".format(core.BASE[t] + core.HEIGHT[t]),
                     "%.0f" % core.HEIGHT[t], "{:,}".format(core.GPF[t]), "{:,}".format(core.nominal(t))])
    return rows


def facilities_docx(path):
    from docx import Document
    from docx.shared import Pt
    d = Document()
    stl = d.styles["Normal"]
    stl.font.name = "Calibri"
    stl.font.size = Pt(10.5)
    for kind, text in FACILITIES:
        if kind == "t":
            d.add_heading(text, level=0)
        elif kind == "s":
            p = d.add_paragraph(text)
            p.runs[0].italic = True
        elif kind == "h":
            d.add_heading(text, level=1)
        elif kind.startswith("TABLE"):
            data = _wells_table() if kind == "TABLE_WELLS" else _tanks_table()
            tb = d.add_table(rows=len(data), cols=len(data[0]))
            tb.style = "Light Grid Accent 1"
            for i, r in enumerate(data):
                for j, v in enumerate(r):
                    cell = tb.cell(i, j)
                    cell.text = v
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(8.5 if kind == "TABLE_TANKS" else 9.5)
                            run.bold = i == 0
            d.add_paragraph("")
        else:
            d.add_paragraph(text)
    cp = d.core_properties
    cp.author = "Shelly McCarthy"
    cp.last_modified_by = "Wayne Hicks"
    cp.title = "Campus Water System: Facilities Summary"
    cp.comments = ""
    cp.keywords = ""
    cp.subject = ""
    cp.category = ""
    cp.revision = 7
    cp.created = dt.datetime(2019, 2, 4, 9, 30)
    cp.modified = dt.datetime(2026, 7, 9, 15, 12)
    buf = io.BytesIO()
    d.save(buf)
    normalize_docx(buf.getvalue(), path)


def facilities_text():
    out = []
    for k, t in FACILITIES:
        if k == "TABLE_WELLS":
            out += [" | ".join(r) for r in _wells_table()]
        elif k == "TABLE_TANKS":
            out += [" | ".join(r) for r in _tanks_table()]
        else:
            out.append(t)
    return "\n".join(out)


def normalize_docx(data, path):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        parts = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    drop = {"word/stylesWithEffects.xml", "docProps/thumbnail.jpeg"}
    drop |= {n for n, _ in parts if n.startswith("customXml/")}
    out = []
    for name, b in parts:
        if name in drop:
            continue
        if name.endswith(".rels") or name == "[Content_Types].xml":
            b = re.sub(rb'<Relationship [^>]*Target="[^"]*(stylesWithEffects\.xml|thumbnail\.jpeg|customXml/[^"]*)"/>',
                       b"", b)
            b = re.sub(rb'<Override PartName="/(word/stylesWithEffects\.xml|docProps/thumbnail\.jpeg|customXml/[^"]*)"'
                       rb'[^>]*/>', b"", b)
            b = re.sub(rb'<Default Extension="jpeg"[^>]*/>', b"", b)
        if name.endswith((".xml", ".rels")):
            b = b.replace(b"<?xml version='1.0' encoding='UTF-8' standalone='yes'?>\n", PP.DECL)
        if name == "docProps/app.xml":
            b = re.sub(rb"<Application>[^<]*</Application>", b"<Application>Microsoft Office Word</Application>", b)
            b = re.sub(rb"<AppVersion>[^<]*</AppVersion>", b"<AppVersion>16.0000</AppVersion>", b)
            b = re.sub(rb"<Company>[^<]*</Company>", b"", b)
        if name == "docProps/core.xml":
            b = re.sub(rb"<cp:lastPrinted>[^<]*</cp:lastPrinted>", b"", b)
            b = b.replace(b"python-docx", b"")
        out.append((name, b))
    with zipfile.ZipFile(path, "w") as z:
        for name, b in out:
            zi = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, b)


README = """Kettle Ridge University, Campus Utilities
Water SCADA historian export

Requested by:  Campus Utilities, for the 2026 sanitary survey response
Exported:      6 October 2026, from the water SCADA historian
File:          {scada}
Tag list:      {tags}

COVERAGE

  Every tag in the tag list, one sample per hour, from 2025-10-01T04:00:00Z to
  2026-10-02T03:00:00Z ({n:,} samples a tag, {rows:,} rows).

COLUMNS

  tag             tag name, as in the tag list
  timestamp_utc   sample time, UTC (ISO 8601)
  value           the sample's value, in the tag's units
  quality         G good
                  U uncertain: communications lost, last good value held
                  B bad: instrument fault

VALUES

  Timestamps are UTC.
  Levels and flows are instantaneous samples at the timestamp; so are the pressures
  and the chlorine residual. Totalizers are cumulative gallons, the meter register
  read at the timestamp. Run status is 1 when the pump is running at the timestamp.

  Rows are in time order, the tags of each hour in tag-list order.
"""


def readme_txt(path, files, n, rows):
    with open(path, "w", newline="\r\n", encoding="utf-8") as f:
        f.write(README.format(scada=files["scada"], tags=files["tags"], n=n, rows=rows))


GUIDE = """Kettle Ridge University, Campus Utilities
Description of the records sent to the district office with the 2026 sanitary survey response
12 October 2026

Mr. Reyes asked at the August visit for the raw records behind the response rather than summaries.
They are the systems' own exports; this note says what each file's rows and columns are. The SCADA
historian export has its own readme.

OPERATORS' DAILY LOG  ({log})
  One sheet a month, one row a date.
  W1 meter (gal) to W5 meter (gal)   well discharge meter totalizer at 00:00 on the date
  W1 hrs to W5 hrs                   pump run time on the date, 00:00 to 24:00, hours
  BPS P1 hrs, BPS P2 hrs             booster pump run time on the date, hours
  Cl2 EP (mg/L)                      entry point free chlorine, morning grab sample
  Remarks                            operator's notes

PUMP TEST RECORDS  ({tests})
  Pump tests: one row a test. Pumping rate is the rate held through the test.
  Wells: one row a well, with its nameplate rate.

AMR MONTH-END READS  ({amr})
  One row an endpoint and month-end read.
  endpoint_id     the radio endpoint on the meter register
  meter_no        the meter the endpoint was on at the read
  read_date       the date of the read; read_time is the time, local (00:00 on the 1st)
  register_read   the register reading, in register units (see the meter inventory)

METER INVENTORY  ({meters})
  Meters: one row a meter in service on or after 1 October 2025.
  Bldg No.             the space inventory building number the meter was set up under
  Register multiplier  register reading x multiplier = gallons
  Initial read         register reading when the meter was set
  Final read           register reading when the meter was removed

CMMS WORK ORDERS  ({cmms})
  One row a work order on the water system, opened 1 October 2025 to 9 October 2026, all closed.
  reading_1, reading_2  register readings recorded on the work order: on a meter read or a check read,
                        reading_1 is the register at reading_for; on a change-out, reading_1 is the removed
                        meter's final read and reading_2 the new meter's initial read; on a removal,
                        reading_1 is the final read; on an install, reading_2 is the initial read
  reading_for           the time the logged register applies to, local

SPACE INVENTORY  ({space})
  Buildings: one row a building in the inventory on 30 September 2026. Primary use is the
  space inventory's classification.
  Removed from inventory: buildings taken out of the inventory since 2023.

SERVICE LINE INVENTORY  ({lsl})
  The service line inventory filed on 15 October 2024: one row a service line from a main to a
  building, with its use, size and materials.
"""


def guide_txt(path, files):
    with open(path, "w", newline="\r\n", encoding="utf-8") as f:
        f.write(GUIDE.format(**files))


INDEX = [
    ("letter", "PA DEP Safe Drinking Water Program", "Sanitary survey letter of 15 September 2026",
     "2026 survey", "received 2026-09-17"),
    ("facilities", "Campus Utilities", "Facilities summary of the campus water system", "Revised July 2026",
     "issued 2026-07-09"),
    ("scada", "Water SCADA historian", "Hourly samples of every water SCADA tag",
     "2025-10-01 to 2026-10-01 (local)", "pulled 2026-10-06"),
    ("tags", "Water SCADA historian", "Tag list", "Current", "pulled 2026-10-06"),
    ("readme", "Water SCADA historian", "Readme for the historian export", "", "pulled 2026-10-06"),
    ("log", "Water operators", "Operators' daily log", "October 2025 to 11 October 2026", "copied 2026-10-12"),
    ("tests", "Water operators", "Well pump test records and well data", "2021 to September 2026",
     "copied 2026-09-21"),
    ("amr", "AMR fixed network head end", "Month-end register reads of the meters on the fixed network",
     "Reads of 1 October 2025 to 1 October 2026", "pulled 2026-10-07"),
    ("meters", "Meter Shop", "Meter inventory", "As of 8 October 2026", "pulled 2026-10-08"),
    ("cmms", "CMMS (Facilities work management)", "Water work orders",
     "Opened 1 October 2025 to 9 October 2026", "pulled 2026-10-09"),
    ("space", "Office of Space Management", "Building inventory", "As of 30 September 2026", "received 2026-10-01"),
    ("lsl", "Campus Utilities", "Service line inventory as filed with the Department", "Filed 15 October 2024",
     "filed 2024-10-15"),
    ("weather", "Department of Meteorology campus station", "Daily weather record",
     "1 October 2025 to 30 September 2026", "received 2026-10-05"),
    ("irrigation", "Grounds Department", "Irrigation controller programs", "2026 season", "received 2026-04-17"),
    ("ccr", "Kettle Ridge Borough Water Authority", "2025 annual drinking water quality report",
     "Calendar 2025", "downloaded 2026-07-01"),
    ("guide", "Campus Utilities", "Field guide to these records", "", "written 2026-10-12"),
]


def index_csv(path, files):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(["file", "originator", "content", "coverage", "date"])
        for k, sysr, content, cov, on in INDEX:
            w.writerow([files[k], sysr, content, cov, on])


CCR_ROWS = [
    ["Contaminant", "Unit", "MCL", "MCLG", "Level found", "Range", "Sample year", "Violation", "Likely source"],
    ["Nitrate", "ppm", "10", "10", "3.1", "1.8 - 3.1", "2025", "No", "Fertilizer runoff"],
    ["Barium", "ppm", "2", "2", "0.08", "0.05 - 0.08", "2024", "No", "Erosion of natural deposits"],
    ["Chlorine", "ppm", "MRDL 4", "MRDLG 4", "1.12", "0.58 - 1.40", "2025", "No", "Water additive"],
    ["TTHM", "ppb", "80", "n/a", "11.4", "6.9 - 14.2", "2025", "No", "By-product of chlorination"],
    ["HAA5", "ppb", "60", "n/a", "4.2", "2.1 - 5.6", "2025", "No", "By-product of chlorination"],
    ["Lead (90th percentile)", "ppb", "AL 15", "0", "3", "ND - 9", "2024", "No", "Household plumbing"],
    ["Copper (90th percentile)", "ppm", "AL 1.3", "1.3", "0.21", "0.02 - 0.44", "2024", "No", "Household plumbing"],
]


def ccr_pdf(path):
    st = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=LETTER, leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                            topMargin=0.8 * inch, bottomMargin=0.8 * inch,
                            title="2025 Annual Drinking Water Quality Report",
                            author="Kettle Ridge Borough Water Authority", creator="Kettle Ridge Borough Water Authority",
                            producer="Kettle Ridge Borough Water Authority")
    story = [Paragraph("Kettle Ridge Borough Water Authority", st["lh"]),
             Paragraph("2025 Annual Drinking Water Quality Report, PWSID 4140208", st["lh2"]),
             Paragraph("We are pleased to report that the water the Authority delivered in 2025 met every federal and "
                       "state drinking water standard. This report covers January 1 to December 31, 2025.",
                       st["body"]),
             Paragraph("Our sources", st["head"]),
             Paragraph("The Authority serves about 9,400 people in the Borough and parts of Ferguson Township from "
                       "three wells and the Borough Spring. Water is disinfected with chlorine and treated with "
                       "orthophosphate for corrosion control. A source water assessment completed in 2019 found the "
                       "wells moderately susceptible to agricultural activity.", st["body"]),
             Paragraph("Detected contaminants", st["head"])]
    tb = Table([[Paragraph(c, st["cell"]) for c in r] for r in CCR_ROWS], repeatRows=1,
               colWidths=[1.25 * inch, 0.45 * inch, 0.6 * inch, 0.6 * inch, 0.6 * inch, 0.8 * inch, 0.55 * inch,
                          0.55 * inch, 1.35 * inch])
    tb.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#8899AA")),
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6EEF5")),
                            ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [tb, Spacer(1, 10),
              Paragraph("Definitions", st["head"]),
              Paragraph("Maximum Contaminant Level (MCL): the highest level of a contaminant allowed in drinking "
                        "water. Maximum Contaminant Level Goal (MCLG): the level below which there is no known or "
                        "expected risk to health. Action Level (AL): the concentration which, if exceeded, triggers "
                        "treatment or other requirements. MRDL: maximum residual disinfectant level. ND: not "
                        "detected.", st["small"]),
              Paragraph("Lead in drinking water", st["head"]),
              Paragraph("If present, elevated levels of lead can cause serious health problems, especially for "
                        "pregnant women and young children. Lead in drinking water is primarily from materials and "
                        "components associated with service lines and home plumbing. The Authority's service line "
                        "inventory is available at the Borough office. When your water has been sitting for several "
                        "hours, you can minimize the potential for lead exposure by flushing your tap for 30 seconds "
                        "to 2 minutes before using water for drinking or cooking.", st["body"]),
              Paragraph("Questions about this report: Kettle Ridge Borough Water Authority, 301 Spring Street, "
                        "weekdays 8:00 a.m. to 4:30 p.m. Board meetings are held on the second Tuesday of each "
                        "month at 7:00 p.m. at the Borough building.", st["body"])]
    doc.build(story, onFirstPage=_footer("Kettle Ridge Borough Water Authority  |  2025 Water Quality Report"),
              onLaterPages=_footer("Kettle Ridge Borough Water Authority  |  2025 Water Quality Report"))
    open(path, "wb").write(buf.getvalue())
    PP.normalize_pdf(path, b"D:20260624141500-04'00'", "Kettle Ridge Borough Water Authority")
