"""The pack's documents: the commission's coverage rule, the crash extract's data dictionary, the Department's travel
monitoring report and crash unit costs (PDF); the road inventory, the operator register and the vehicle registrations
(XLSX); and the pack's contents (TXT)."""
from __future__ import annotations

import datetime as dt

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                    # noqa: E402
from reportlab.lib.pagesizes import letter                                          # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                     # noqa: E402
from reportlab.lib.units import inch                                                # noqa: E402
from reportlab.pdfbase import pdfmetrics                                            # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                        # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle   # noqa: E402

import common as K   # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("Sans", "LiberationSans-Regular"), ("Sans-B", "LiberationSans-Bold"), ("Sans-I", "LiberationSans-Italic"),
              ("Serif", "LiberationSerif-Regular"), ("Serif-B", "LiberationSerif-Bold")):
    pdfmetrics.registerFont(TTFont(nm, FONT_DIR + f + ".ttf"))
pdfmetrics.registerFontFamily("Sans", normal="Sans", bold="Sans-B", italic="Sans-I", boldItalic="Sans-B")
pdfmetrics.registerFontFamily("Serif", normal="Serif", bold="Serif-B", italic="Serif", boldItalic="Serif-B")
INK, GREY, RULE = colors.HexColor("#1c2329"), colors.HexColor("#5a646d"), colors.HexColor("#b9c1c8")
GREEN = colors.HexColor("#1f4d3a")
BLUE = colors.HexColor("#22406b")


def styles(face="Sans", accent=BLUE):
    b = face + "-B"
    body = ParagraphStyle("b", fontName=face, fontSize=9.6, leading=13.2, textColor=INK, spaceAfter=5)
    return dict(
        body=body,
        h1=ParagraphStyle("h1", fontName=b, fontSize=14.5, leading=18, textColor=accent, spaceAfter=2),
        h2=ParagraphStyle("h2", fontName=b, fontSize=10.4, leading=14, textColor=INK, spaceBefore=8, spaceAfter=2,
                          keepWithNext=1),
        meta=ParagraphStyle("m", fontName=face, fontSize=8.3, leading=11.2, textColor=GREY, spaceAfter=9),
        cell=ParagraphStyle("c", fontName="Sans", fontSize=7.8, leading=9.8, textColor=INK),
        cellb=ParagraphStyle("cb", fontName="Sans-B", fontSize=7.8, leading=9.8, textColor=INK),
        note=ParagraphStyle("n", fontName="Sans", fontSize=7.6, leading=9.8, textColor=GREY, spaceBefore=2,
                            spaceAfter=6),
        clause=ParagraphStyle("p", parent=body, leftIndent=26, firstLineIndent=-26))


def _frame(left, right):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(GREY)
        canvas.setFont("Sans", 7.4)
        canvas.drawString(.8 * inch, letter[1] - .5 * inch, left)
        canvas.drawRightString(letter[0] - .8 * inch, letter[1] - .5 * inch, right)
        canvas.drawRightString(letter[0] - .8 * inch, .45 * inch, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _table(S, rows, widths, right=(), bold_last=False):
    data = [[Paragraph(str(c), S["cellb"] if (i == 0 or (bold_last and i == len(rows) - 1)) else S["cell"])
             for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=[w * inch for w in widths], repeatRows=1, hAlign="LEFT")
    st = [("LINEBELOW", (0, 0), (-1, 0), .6, INK), ("LINEBELOW", (0, -1), (-1, -1), .4, RULE),
          ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 1.6),
          ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6)]
    for c in right:
        st.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
    t.setStyle(TableStyle(st))
    return t


def _clause(S, num, text):
    return Paragraph("<b>%s</b>&nbsp;&nbsp;%s" % (num, text), S["clause"])


def _doc(path, title, author, subject, left, right):
    return SimpleDocTemplate(path, pagesize=letter, leftMargin=.8 * inch, rightMargin=.8 * inch,
                             topMargin=.8 * inch, bottomMargin=.75 * inch, title=title, author=author,
                             subject=subject, creator=author, producer=author), _frame(left, right)


# ------------------------------------------------------------------------------------------------ the coverage rule
def write_rule(path):
    S = styles("Serif", GREEN)
    doc, fr = _doc(path, "Commercial Fleet Safety Programme: Coverage Rule", K.COMMISSION, "Coverage rule",
                   K.COMMISSION, "Commercial Fleet Safety Programme: Coverage Rule")
    st = [Paragraph("Commercial Fleet Safety Programme: Coverage Rule", S["h1"]),
          Paragraph("%s. Adopted 14 October 2026 (Resolution 2026-41). Rule owner: %s, Director of Safety "
                    "Programmes." % (K.COMMISSION, K.OWNER), S["meta"]),
          Paragraph("1. Purpose", S["h2"]),
          Paragraph("From 1 July 2027 every operator of a covered use fits each vehicle it runs in that use with a "
                    "telematics unit and enrols the drivers of the use in the programme's coaching. Operators pay for "
                    "both. The Commission covers the uses whose drivers cause crashes most often for the distance "
                    "they drive.", S["body"]),
          Paragraph("2. The uses", S["h2"]),
          Paragraph("The Commission licenses seven commercial uses of motor vehicles in the region:", S["body"])]
    desc = {
        "TX": "Vehicles for hire of a licensed taxicab company, dispatched or hailed",
        "RH": "Personal vehicles carrying passengers on trips arranged through a licensed transportation network "
              "company's app",
        "LM": "Pre-arranged passenger service by a licensed limousine or chauffeured car operator",
        "AS": "Scheduled and on-demand shuttles of a licensed shuttle operator serving the airport and hotels",
        "PT": "Licensed operators carrying riders with disabilities or riders to medical appointments",
        "PD": "Vans and trucks of licensed parcel and courier operators",
        "AD": "Couriers' vehicles delivering orders arranged through a licensed delivery network company's app",
    }
    rows = [["Use", "What it covers"]] + [[K.RULE_NAME[c], desc[c]] for c in K.CODES]
    st.append(_table(S, rows, [2.35, 4.45]))
    st += [Paragraph("3. Coverage test", S["h2"]),
           _clause(S, "3.1", "A use is covered if its drivers caused crashes at a rate per mile driven of at least "
                             "1.5 times the rate of private drivers, over 1 January 2021 to 31 December 2025."),
           _clause(S, "3.2", "Each rate is the number of crashes the drivers caused in those five years divided by "
                             "the miles their vehicles drove in the region in the same five years."),
           Paragraph("4. Terms", S["h2"]),
           _clause(S, "4.1", "<i>Crash.</i> A crash on a public road in %s reported by a law enforcement agency to "
                             "the state crash records system, in which two or more motor vehicles in transport were "
                             "involved. Crashes in which a single vehicle in transport left the road or collided with "
                             "an object, an animal, a parked vehicle, a pedestrian or a cyclist are not counted: they "
                             "are reported much less completely than crashes between moving vehicles." % K.COUNTY),
           _clause(S, "4.2", "<i>Caused.</i> A crash is caused by each driver the investigating officer recorded as at "
                             "fault on the crash report: on the 2016 report form, a driver with a contributing "
                             "circumstance recorded; on the 2023 form, the driver of the at-fault unit. A crash with "
                             "no driver recorded at fault is caused by no one. A driver who left the scene and was not "
                             "identified is not a driver of any use."),
           _clause(S, "4.3", "<i>Drivers of a use.</i> The drivers of vehicles operating in the use when the crash "
                             "happened, as the crash report records the vehicle's use and its operator."),
           _clause(S, "4.4", "<i>Private drivers.</i> Drivers of passenger cars, sport utility vehicles, pickups and "
                             "vans in personal use."),
           _clause(S, "4.5", "<i>Miles driven.</i> The vehicle miles travelled by the vehicles concerned on public "
                             "roads in %s." % K.COUNTY),
           Paragraph("5. Decision and review", S["h2"]),
           Paragraph("Commission staff report each use's rate against private drivers' rate and the uses the test "
                     "covers. The Commission tells the operators of each covered use by 15 January 2027, so that "
                     "they can fit their vehicles and enrol their drivers before coverage starts on 1 July 2027. The "
                     "test is run again every three years on the latest five years of crashes.",
                     S["body"])]
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ------------------------------------------------------------------------------------------------ the data dictionary
CRASH_FIELDS = [
    ("CRASH_ID", "State crash number, assigned when the report is received (year and sequence)."),
    ("FORM_VERSION", "Report form: PTCR-16 (crashes to 31 December 2022) or PTCR-23 (from 1 January 2023)."),
    ("AGENCY", "Reporting agency: SP6 State Patrol, District 6; SCSO Stonereach County Sheriff's Office; KPD "
               "Kelderwick Police Department; FPD Fairhollow Police Department."),
    ("AGENCY_REPORT_NO", "The agency's own report number."),
    ("CRASH_DATE, CRASH_TIME", "Date and time of the crash as the agency submitted them. Formats differ by agency: "
                               "HH:MM (SP6, SCSO), h:mm AM/PM (KPD), HHMM with no separator or leading zeros (FPD). "
                               "Blank where no time was recorded."),
    ("TIME_ZONE", "UTC: SP6 submits the date and time in Coordinated Universal Time. PT: Pacific local time (the "
                  "other agencies)."),
    ("ROUTE_ID", "Route of the crash location in the current linear referencing system (PTCR-23 crashes)."),
    ("LOC_ROUTE", "Route number as written on PTCR-16 reports; the road inventory's sheet Legacy route numbers gives "
                  "the current ROUTE_ID."),
    ("MILEPOST", "Milepost of the crash location on the route."),
    ("F_SYSTEM, URBAN_RURAL", "Functional system (1 to 7) and urban (U) or rural (R) area of the location, added by "
                              "location coding of PTCR-23 crashes. Blank for PTCR-16 crashes."),
    ("CITY", "City whose limits contain the location; blank outside city limits."),
    ("FIRST_HARMFUL_EVENT", "The first injury- or damage-producing event (codes below)."),
    ("MANNER_OF_COLLISION", "For a collision between motor vehicles in transport (codes below)."),
    ("NUM_UNITS", "Units on the report: motor vehicles in transport, parked vehicles, pedestrians and "
                  "pedalcyclists."),
    ("AT_FAULT_UNIT", "PTCR-23: the unit the investigating officer recorded at fault; 0 when none was recorded."),
    ("HIT_AND_RUN", "Y when a unit left the scene."),
    ("LIGHT_CONDITION, WEATHER, ROAD_SURFACE", "As recorded."),
    ("CRASH_SEVERITY", "The most severe injury recorded on the crash report (KABCO; PTCR-16 codes 1 to 5)."),
    ("NUM_KILLED, NUM_INJURED", "Persons recorded killed, and injured (A, B or C), on the crash report."),
]
UNIT_FIELDS = [
    ("CRASH_ID, UNIT_NO", "The crash and the unit's number on the report."),
    ("UNIT_TYPE", "PTCR-16: 1 motor vehicle in transport, 2 parked motor vehicle, 3 pedestrian, 4 pedalcyclist. "
                  "PTCR-23: MV IN TRANSPORT, PARKED MV, PEDESTRIAN, PEDALCYCLIST."),
    ("BODY_TYPE", "Vehicle body (codes below)."),
    ("VEH_USE", "The use the vehicle was operating in (codes below)."),
    ("CARRIER_NAME", "Operator, carrier or owning agency as the officer recorded it, for vehicles in other than "
                     "personal use."),
    ("CONTRIB_CODE", "PTCR-16: contributing circumstance recorded for the driver or non-motorist; 00 none recorded. "
                     "Blank on PTCR-23 reports."),
    ("HIT_RUN_UNIT", "Y when the unit left the scene."),
    ("OCCUPANTS", "Persons in the vehicle, driver included."),
    ("DRIVER_AGE, MODEL_YEAR", "As recorded."),
    ("STRIKING_UNIT", "PTCR-23, pedestrians and pedalcyclists: the vehicle unit in collision with them."),
]
PERSON_FIELDS = [
    ("CRASH_ID, UNIT_NO, PERSON_NO", "The crash, the unit the person was in or was (for a non-motorist), and the "
                                     "person's number."),
    ("PERSON_TYPE", "PTCR-16: 1 driver, 2 passenger, 3 pedestrian, 4 pedalcyclist. PTCR-23: DRIVER, PASSENGER, "
                    "PEDESTRIAN, BICYCLIST."),
    ("SEAT_POSITION", "01 driver's seat; 02 to 12 the other seating positions; 00 not in a vehicle (pedestrians "
                      "and pedalcyclists)."),
    ("AGE, SEX", "As recorded."),
    ("INJURY_SEVERITY", "KABCO as recorded (PTCR-16 codes 1 to 5)."),
]
CODES_TABLES = [
    ("VEH_USE", ["Code (PTCR-16)", "Code (PTCR-23)", "Use"], [
        ["0", "00", "Personal"], ["1", "11", "For hire, passenger: taxicab (PTCR-23 11)"],
        ["1", "12", "For hire, passenger: ride-hail, transportation network company (PTCR-23 12)"],
        ["1", "13", "For hire, passenger: limousine or chauffeured car (PTCR-23 13)"],
        ["1", "14", "For hire, passenger: shuttle (PTCR-23 14)"],
        ["1", "15", "For hire, passenger: paratransit or non-emergency medical (PTCR-23 15)"],
        ["2", "21", "Delivery: parcel or courier (PTCR-23 21)"],
        ["2", "22", "Delivery: app-based, delivery network company (PTCR-23 22)"],
        ["3", "30", "Other commercial"], ["4", "41", "Bus: transit (PTCR-23 41)"],
        ["4", "42", "Bus: school (PTCR-23 42)"], ["5", "50", "Government"], ["6", "60", "Emergency"],
        ["9", "99", "Unknown"]]),
    ("BODY_TYPE", ["Code (PTCR-16)", "Code (PTCR-23)", "Body"], [
        ["01", "PASSENGER CAR", "Passenger car"], ["02", "SPORT UTILITY", "Sport utility vehicle"],
        ["03", "PICKUP", "Pickup"], ["04", "VAN", "Passenger van or minivan"], ["05", "CARGO VAN", "Cargo van"],
        ["06", "MINIBUS", "Minibus or shuttle bus (9 to 15 seats)"], ["07", "BUS", "Bus (16 or more seats)"],
        ["08", "SINGLE-UNIT TRUCK", "Single-unit truck"], ["09", "TRUCK-TRACTOR", "Truck-tractor"],
        ["10", "MOTORCYCLE", "Motorcycle"], ["99", "UNKNOWN", "Unknown"]]),
    ("FIRST_HARMFUL_EVENT", ["Code (PTCR-16)", "Code (PTCR-23)", "Event"], [
        ["01", "MOTOR VEHICLE IN TRANSPORT", "Collision with a motor vehicle in transport"],
        ["02", "PARKED MOTOR VEHICLE", "Collision with a parked motor vehicle"],
        ["03", "PEDESTRIAN", "Collision with a pedestrian"], ["04", "PEDALCYCLIST", "Collision with a pedalcyclist"],
        ["05", "FIXED OBJECT", "Collision with a fixed object"], ["06", "OVERTURN/ROLLOVER", "Overturn"],
        ["07", "ANIMAL", "Collision with an animal"], ["08", "OTHER OBJECT", "Other object or non-collision"]]),
    ("MANNER_OF_COLLISION", ["Code (PTCR-16)", "Code (PTCR-23)", "Manner"], [
        ["1", "REAR END", "Front to rear"], ["2", "ANGLE", "Angle"], ["3", "SIDESWIPE SAME DIRECTION",
                                                                     "Sideswipe, same direction"],
        ["4", "SIDESWIPE OPPOSITE", "Sideswipe, opposite direction"], ["5", "HEAD ON", "Front to front"],
        ["6", "OTHER", "Other"]]),
    ("CONTRIB_CODE (PTCR-16)", ["Code", "", "Contributing circumstance"], [
        ["00", "", "None recorded"], ["01", "", "Failed to yield right of way"], ["02", "", "Followed too closely"],
        ["03", "", "Exceeded safe speed"], ["04", "", "Disregarded traffic signal or sign"],
        ["05", "", "Improper lane change"], ["06", "", "Inattention or distraction"], ["07", "", "Improper turn"],
        ["08", "", "Under the influence"], ["09", "", "Improper backing"], ["12", "", "Wrong side or wrong way"],
        ["14", "", "Other improper driving"], ["31", "", "Non-motorist: entered roadway improperly"],
        ["32", "", "Non-motorist: not in crosswalk"], ["33", "", "Non-motorist: against signal"],
        ["34", "", "Non-motorist: other improper action"]]),
    ("INJURY_SEVERITY, CRASH_SEVERITY", ["Code (PTCR-16)", "Code (PTCR-23)", "Severity"], [
        ["1", "K", "Fatal injury"], ["2", "A", "Suspected serious injury"], ["3", "B", "Suspected minor injury"],
        ["4", "C", "Possible injury"], ["5", "O", "No apparent injury"]]),
]


def write_dictionary(path):
    S = styles("Sans", BLUE)
    doc, fr = _doc(path, "State Crash Records System: Regional Extract, Stonereach County", "State Crash Records Unit",
                   "Data dictionary", "State Crash Records System", "Regional extract: data dictionary")
    st = [Paragraph("State Crash Records System: Regional Extract, Stonereach County", S["h1"]),
          Paragraph("Data dictionary. Extract of 30 September 2026, prepared for the %s. Crash Records Unit, "
                    "Traffic Records Section." % K.COMMISSION, S["meta"]),
          Paragraph("The extract", S["h2"]),
          Paragraph("Crashes reported by law enforcement agencies on public roads in Stonereach County that "
                    "happened from 1 January 2021 to 31 December 2025, in three files: crashes (one row a crash), "
                    "units (one row a unit on the report: a motor vehicle in transport, a parked vehicle, a "
                    "pedestrian or a pedalcyclist) and persons (the people involved in crashes with an injury "
                    "recorded). A crash is reportable when someone is injured or killed or the "
                    "damage is $1,000 or more.", S["body"]),
          Paragraph("Two editions of the Police Traffic Crash Report were in use: the 2016 edition (PTCR-16) for "
                    "crashes to 31 December 2022 and the 2023 edition (PTCR-23) from 1 January 2023. Field contents "
                    "and codes differ between them as shown.", S["body"]),
          Paragraph("Crashes file", S["h2"]),
          _table(S, [["Field", "Contents"]] + [list(r) for r in CRASH_FIELDS], [1.7, 5.1]),
          Paragraph("Units file", S["h2"]),
          _table(S, [["Field", "Contents"]] + [list(r) for r in UNIT_FIELDS], [1.7, 5.1]),
          Paragraph("Persons file", S["h2"]),
          _table(S, [["Field", "Contents"]] + [list(r) for r in PERSON_FIELDS], [1.7, 5.1]),
          Paragraph("Codes", S["h2"])]
    for name, head, rows in CODES_TABLES:
        st.append(Paragraph(name, S["cellb"]))
        st.append(_table(S, [head] + rows, [1.1, 2.1, 3.6]))
        st.append(Spacer(1, 6))
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ------------------------------------------------------------------------------------------------ travel monitoring
def write_tmr(path, vmt, daily, stations, monthly, classes):
    """vmt: {(fc, ur): {year: millions}}; daily: {(fc, ur): (centerline miles, daily VMT thousands)};
    stations: rows; monthly: {group: [12 factors]}; classes: [(class, share)]."""
    S = styles("Sans", BLUE)
    doc, fr = _doc(path, "Travel Monitoring Report 2025", K.DOT, "Travel monitoring", "%s: Transportation Data Office"
                   % K.DOT, "Travel Monitoring Report 2025")
    fsys = {1: "Interstate", 2: "Other freeways and expressways", 3: "Other principal arterials",
            4: "Minor arterials", 5: "Major collectors", 6: "Minor collectors", 7: "Local"}
    st = [Paragraph("Travel Monitoring Report 2025", S["h1"]),
          Paragraph("%s, Transportation Data Office. Stonereach Valley region (Stonereach County). Issued 18 August "
                    "2026." % K.DOT, S["meta"]),
          Paragraph("1. Vehicle miles of travel", S["h2"]),
          Paragraph("Vehicle miles of travel (VMT) on public roads are estimated each year from the traffic counts on "
                    "each road section and its length, as reported to the Federal Highway Administration. Table 1 "
                    "gives annual VMT for the region by functional system and area; Table 2 gives 2025 daily VMT and "
                    "centerline miles.", S["body"]),
          Paragraph("Table 1. Annual vehicle miles of travel, Stonereach Valley region (millions)", S["cellb"])]
    rows = [["Functional system", "Area"] + [str(y) for y in K.YEARS]]
    tot = {y: 0.0 for y in K.YEARS}
    for fc, ur, g in K.FCLASS:
        r = [fsys[fc], "Urban" if ur == "U" else "Rural"]
        for y in K.YEARS:
            r.append("{:,.1f}".format(vmt[(fc, ur)][y]))
            tot[y] += vmt[(fc, ur)][y]
        rows.append(r)
    rows.append(["All roads", ""] + ["{:,.1f}".format(tot[y]) for y in K.YEARS])
    st.append(_table(S, rows, [1.9, 0.7] + [0.84] * 5, right=(2, 3, 4, 5, 6), bold_last=True))
    st.append(Paragraph("Totals are of the unrounded estimates.", S["note"]))
    st.append(Paragraph("Table 2. Centerline miles and daily vehicle miles of travel, 2025", S["cellb"]))
    rows = [["Functional system", "Area", "Centerline miles", "Daily VMT (thousands)"]]
    for fc, ur, g in K.FCLASS:
        cm, dv = daily[(fc, ur)]
        rows.append([fsys[fc], "Urban" if ur == "U" else "Rural", "{:,.1f}".format(cm), "{:,.1f}".format(dv)])
    st.append(_table(S, rows, [1.9, 0.7, 1.2, 1.5], right=(2, 3)))
    st += [Paragraph("2. Continuous count program", S["h2"]),
           Paragraph("Continuous count stations record the traffic passing in each direction every hour of the year. "
                     "Each station belongs to the factor group of its road (Table 3). The stations' 2025 hourly volumes "
                     "are published with this report (ccs_hourly_volumes_2025.csv, one row a station, day and "
                     "direction; HR01 is the hour from midnight to 01:00 local time, HR24 the hour to midnight). Days "
                     "a station did not record are left out.", S["body"]),
           Paragraph("Table 3. Factor groups", S["cellb"])]
    rows = [["Group", "Name", "Functional systems"]]
    fs_of = {"UF": "1, 2 (urban)", "UA": "3, 4 (urban)", "UL": "5, 7 (urban)", "RF": "1, 2 (rural)",
             "RA": "3, 4 (rural)", "RL": "5, 6, 7 (rural)"}
    for i, g in enumerate(K.GROUPS, 1):
        rows.append([str(i), K.GROUP_NAME[g], fs_of[g]])
    st.append(_table(S, rows, [0.6, 3.4, 2.0]))
    st.append(Spacer(1, 6))
    st.append(Paragraph("Table 4. Continuous count stations in the region, 2025", S["cellb"]))
    st.append(_table(S, [["Station", "Route", "Milepost", "Functional system", "Area", "Group"]] + stations,
                     [0.9, 0.9, 0.8, 1.2, 0.6, 0.6], right=(2,)))
    st.append(Spacer(1, 6))
    st.append(Paragraph("Table 5. Monthly factors, 2025 (average day of the month against the annual average day)",
                        S["cellb"]))
    rows = [["Group"] + [dt.date(2025, m, 1).strftime("%b") for m in range(1, 13)]]
    for i, g in enumerate(K.GROUPS, 1):
        rows.append([str(i)] + ["%.2f" % x for x in monthly[g]])
    st.append(_table(S, rows, [0.6] + [0.5] * 12, right=tuple(range(1, 13))))
    st += [Paragraph("3. Travel by vehicle class", S["h2"]),
           Paragraph("Classification counts give the share of VMT by vehicle class (FHWA scheme, grouped).",
                     S["body"]),
           Paragraph("Table 6. Share of vehicle miles of travel by vehicle class, 2025", S["cellb"]),
           _table(S, [["Vehicle class", "Share of VMT"]] + [[c, "%.1f%%" % (100 * s)] for c, s in classes],
                  [3.0, 1.2], right=(1,))]
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ------------------------------------------------------------------------------------------------ crash unit costs
COSTS = {"K": 12850000, "A": 1183000, "B": 318000, "C": 161000, "O": 15300}


def write_costs(path):
    S = styles("Sans", BLUE)
    doc, fr = _doc(path, "Crash Unit Costs for Safety Analysis, 2025", K.DOT, "Crash unit costs",
                   "%s: Office of Traffic Safety" % K.DOT, "Crash unit costs, 2025")
    st = [Paragraph("Crash Unit Costs for Safety Analysis, 2025", S["h1"]),
          Paragraph("%s, Office of Traffic Safety. Highway Safety Improvement Program guidance note 2025-03, issued "
                    "6 March 2025." % K.DOT, S["meta"]),
          Paragraph("Benefit-cost analyses for safety projects, and reports of the cost of crashes, use the "
                    "comprehensive crash unit costs below. They value the economic costs of a crash (medical, "
                    "emergency services, lost productivity, property damage, insurance and legal) and the lost quality "
                    "of life, in 2025 dollars.", S["body"]),
          Paragraph("A crash is costed once, at the unit cost of its most severe injury on the KABCO scale. A fatal "
                    "crash (K) is one in which a person died within 30 days of the crash.", S["body"]),
          _table(S, [["Crash severity", "Unit cost a crash (2025 dollars)"],
                     ["K  Fatal", "{:,}".format(COSTS["K"])], ["A  Suspected serious injury", "{:,}".format(COSTS["A"])],
                     ["B  Suspected minor injury", "{:,}".format(COSTS["B"])],
                     ["C  Possible injury", "{:,}".format(COSTS["C"])],
                     ["O  Property damage only", "{:,}".format(COSTS["O"])]], [2.6, 2.2], right=(1,)),
          Paragraph("Source: the Department's crash cost study (2024), updated to 2025 dollars with the consumer price "
                    "index and the employment cost index.", S["note"])]
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ------------------------------------------------------------------------------------------------ workbooks
def _style_sheet(sh, widths=None, fill="22406B"):
    from openpyxl.styles import Alignment, Font, PatternFill
    for cell in sh[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(vertical="center")
    sh.freeze_panes = "A2"
    for col in sh.columns:
        letter_ = col[0].column_letter
        if widths and letter_ in widths:
            sh.column_dimensions[letter_].width = widths[letter_]
        else:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:400])
            sh.column_dimensions[letter_].width = min(48, max(9, width + 2))


def write_inventory(path, segs, legacy):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Segments"
    ws.append(["ROUTE_ID", "BEGIN_MP", "END_MP", "F_SYSTEM", "URBAN_RURAL", "CITY", "AADT_2025", "THROUGH_LANES"])
    for s in sorted(segs, key=lambda s: (s["route"], s["b"])):
        lanes = {1: 6 if s["ur"] == "U" else 4, 2: 4, 3: 4 if s["ur"] == "U" else 2, 4: 2}.get(s["fc"], 2)
        ws.append([s["route"], s["b"], s["e"], s["fc"], s["ur"], s["city"].upper() or None, s["aadt"], lanes])
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "0.000"
        row[2].number_format = "0.000"
        row[6].number_format = "#,##0"
    _style_sheet(ws)
    ws.auto_filter.ref = "A1:H%d" % ws.max_row
    w2 = wb.create_sheet("Legacy route numbers")
    w2.append(["LEGACY_ROUTE_NO", "ROUTE_ID"])
    for old, new in legacy:
        w2.append([old, new])
    _style_sheet(w2)
    w3 = wb.create_sheet("Read me")
    w3.append(["Road inventory, Stonereach County"])
    for line in ["Linear referencing system extract, 31 December 2025. One row a section of a route with uniform "
                 "functional system and area: from BEGIN_MP up to, not including, END_MP.",
                 "F_SYSTEM: 1 Interstate, 2 other freeways and expressways, 3 other principal arterials, 4 minor "
                 "arterials, 5 major collectors, 6 minor collectors, 7 local. URBAN_RURAL: U urban area, R rural.",
                 "AADT_2025: annual average daily traffic, both directions.",
                 "Route identifiers changed with the 2023 linear referencing system; Legacy route numbers gives the "
                 "number reports used before."]:
        w3.append([line])
    w3.column_dimensions["A"].width = 120
    wb.save(path)


def write_operators(path, ops, fleet):
    """ops: rows of the operator list; fleet: rows (operator id, name, use, permit year, vehicles)."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Operators"
    ws.append(["OPERATOR_ID", "LICENSED_NAME", "USE", "LICENCE_NO", "LICENSED_FROM", "LICENSED_TO",
               "NAMES_ON_CRASH_REPORTS"])
    for r in ops:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        row[4].number_format = "yyyy-mm-dd"
        row[5].number_format = "yyyy-mm-dd"
    _style_sheet(ws, {"G": 60}, fill="1F4D3A")
    w2 = wb.create_sheet("Fleet returns")
    w2.append(["OPERATOR_ID", "LICENSED_NAME", "USE", "PERMIT_YEAR", "VEHICLES_1_JULY"])
    for r in fleet:
        w2.append(r)
    _style_sheet(w2, fill="1F4D3A")
    wb.save(path)


def write_registrations(path, blocks):
    """blocks: {year: [(vehicle type, private, commercial, for-hire, government)]} as year blocks on one sheet."""
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Registrations"
    ws.append(["Vehicle registrations in force, Stonereach County"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Department of Motor Vehicles, Research and Statistics. Registrations in force on 31 December."])
    for y in sorted(blocks):
        ws.append([])
        ws.append(["%d" % y])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.append(["Vehicle type", "Private plates", "Commercial plates", "For-hire plates", "Government plates",
                   "Total"])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
        for name, a, b, c, d in blocks[y]:
            ws.append([name, a, b, c, d, a + b + c + d])
            for cell in ws[ws.max_row][1:]:
                cell.number_format = "#,##0"
    ws.column_dimensions["A"].width = 34
    for col in "BCDEF":
        ws.column_dimensions[col].width = 16
    wb.save(path)


def write_contents(path, items):
    out = ["%s, Commercial Fleet Safety Programme: coverage review pack" % K.COMMISSION,
           "Assembled %s by the Safety Programmes office." % K.ASSEMBLED.strftime("%d %B %Y").lstrip("0"), ""]
    for name, desc in items:
        out.append(name)
        out.append("    " + desc)
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(out) + "\n")
