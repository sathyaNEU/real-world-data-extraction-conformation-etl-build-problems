"""task273 architecture 5: the golden, computed from the cut pack alone (the files under target/), and the ladder's
routes on the same inputs. Thresholds come from the framework's table (passed in; verify.py reads them off the PDF)."""
from __future__ import annotations

import csv
import json
import os
from datetime import date, datetime, timedelta

import numpy as np

import params as P

HRS = np.arange(P.WINDOW[0], P.WINDOW[1], 0.5)          # half-hours starting 10.00 to 2.30 pm
CALL_FILES = ["spine", "sample", "outlook", "models", "meters", "reads", "lc", "framework", "notice"]
ASK_FILES = ["der", "meters", "framework"]
NAME = {s[0]: s[1] for s in P.SUBS}
CODES = [s[0] for s in P.SUBS]


def rows(path):
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def xlsx_rows(path, sheet=None, header_row=1):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    for _ in range(header_row - 1):
        next(it)
    head = [str(h) for h in next(it)]
    out = [dict(zip(head, list(r) + [None] * (len(head) - len(r)))) for r in it if any(v is not None for v in r)]
    wb.close()
    return out


def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return date.fromisoformat(str(v)[:10])


def local_of_utc(s):
    u = datetime.strptime(s, "%Y-%m-%d %H:%M:%S")
    start = u - timedelta(minutes=30)
    off = 11 if start >= P.DST_START_UTC else 10
    return start + timedelta(hours=off)


def code_of_tag(tag):
    head = tag.split("_")[0]
    if head in NAME:
        return head
    for c, n in NAME.items():
        if n.upper().startswith(head):
            return c
    raise KeyError(tag)


class Pack:
    def __init__(self, target, files):
        self.t = target
        self.f = files

    def p(self, k):
        return os.path.join(self.t, self.f[k])

    def load(self, partial=False):
        # spine: local half-hour starts on each local date
        dem = {}
        for r in rows(self.p("spine")):
            ls = local_of_utc(r["interval_end_utc"])
            dem[(code_of_tag(r["point_name"]), ls)] = float(r["value_mw"])
        self.dem = dem
        js = json.load(open(self.p("sample")))
        pv = {}
        for v in js["values"]:
            end = datetime.strptime(v["interval_end"][:19], "%Y-%m-%dT%H:%M:%S")
            pv[end - timedelta(minutes=30)] = v["output_kw_per_kw"]
        self.pv = pv
        self.outlook = {} if partial else {r["Code"]: float(r["Forecast new inverter capacity (kW)"])
                                           for r in xlsx_rows(self.p("outlook"), header_row=3)}
        self.models = {}
        for r in xlsx_rows(self.p("models")):
            regs = {}
            for part in str(r["Registers"]).split(";"):
                rid, _, what = part.strip().partition(" ")
                regs[what] = rid
            desc = r["Description"]
            cls = "R" if "audio-frequency receiver" in desc else ("T" if "time switch" in desc else "-")
            self.models[r["Model code"]] = dict(cl=regs.get("controlled load") or regs.get("controlled load import"),
                                                cls=cls, data=r["Data"].lower())
        self.meters = rows(self.p("meters"))
        self.reads = rows(self.p("reads"))
        self.lc = rows(self.p("lc"))
        return self


def reference_day(pk):
    tot = {}
    for (c, ls), v in pk.dem.items():
        if P.WINDOW[0] <= ls.hour + ls.minute / 60 < P.WINDOW[1]:
            tot[ls] = tot.get(ls, 0.0) + v
    by_day = {}
    for ls, v in tot.items():
        by_day[ls.date()] = min(by_day.get(ls.date(), 1e9), v)
    sundays = {d: v for d, v in by_day.items() if d.weekday() == 6}
    ref = min(sundays, key=sundays.get)
    return ref, sundays, by_day


def heaters(pk):
    """Per controlled-load heater: substation, group, element kW, daily energy, and the meters' switching."""
    serial_model = {m["meter_serial"]: m["model_code"] for m in pk.meters}
    cur = {}
    first = {}
    loc = {}
    for m in sorted(pk.meters, key=lambda m: m["installed"]):
        first.setdefault(m["nmi"], m["model_code"])
        loc[m["nmi"]] = m["zone_substation"]
        if not m["removed"]:
            cur[m["nmi"]] = m["model_code"]
    cur_serial = {m["nmi"]: m["meter_serial"] for m in pk.meters if not m["removed"]}
    e_tot, e_cur = {}, {}
    for r in pk.reads:
        mdl = pk.models[serial_model[r["meter_serial"]]]
        if r["register_id"] != mdl["cl"]:
            continue
        a = e_tot.setdefault(r["nmi"], [0.0, 0])
        a[0] += float(r["kwh"])
        a[1] += int(r["days"])
        if r["meter_serial"] == cur_serial.get(r["nmi"]):
            e_cur[r["nmi"]] = float(r["kwh"]) / int(r["days"])
    out = []
    for r in pk.lc:
        n = r["nmi"]
        kwh, days = e_tot[n]
        m = pk.models[cur[n]]
        out.append(dict(nmi=n, sub=loc[n], group=int(r["switching_group"]) - 1, kw=float(r["element_kw"]),
                        e=kwh / days, days=days, e_cur=e_cur.get(n, kwh / days), cls=m["cls"], data=m["data"],
                        first_cls=pk.models[first[n]]["cls"]))
    return out


def load(hs, keep, mode="replay", e_key="e"):
    out = {c: np.zeros(len(HRS)) for c in CODES}
    for x, k in zip(hs, keep):
        if not k:
            continue
        r0 = P.SOAK_RELEASE[x["group"]]
        if mode == "replay":
            dur = min(x[e_key] / x["kw"], P.SOAK_OFF - r0)
            out[x["sub"]] += x["kw"] * np.clip(np.minimum(r0 + dur, HRS + 0.5) - np.maximum(r0, HRS), 0, 0.5) / 0.5
        elif mode == "halfhour":
            dur = min(round(x[e_key] / x["kw"] * 2) / 2.0, P.SOAK_OFF - r0)
            out[x["sub"]] += x["kw"] * ((HRS >= r0) & (HRS < r0 + dur))
        else:
            out[x["sub"]] += x[e_key] / (P.SOAK_OFF - P.SOAK_RELEASE[0])
    return out


def routes(hs):
    cls = np.array([x["cls"] for x in hs])
    data = np.array([x["data"] for x in hs])
    fcl = np.array([x["first_cls"] for x in hs])
    n = len(hs)
    out = {"R0": load(hs, [False] * n), "R1": load(hs, [True] * n, "even"), "R2": load(hs, [True] * n),
           "R3": load(hs, cls == "R", "even"), "R4": load(hs, data == "interval"), "R6": load(hs, fcl == "R"),
           "R7": load(hs, cls == "R", e_key="e_cur"), "R8": load(hs, cls == "R", "halfhour"),
           "G": load(hs, cls == "R")}
    share = float((cls == "R").mean())
    out["R5"] = {c: v * share for c, v in out["R2"].items()}
    return out


def window_starts(day):
    return [datetime(day.year, day.month, day.day) + timedelta(hours=float(h)) for h in HRS]


def window_demand(pk, day):
    return {c: np.array([pk.dem[(c, s)] for s in window_starts(day)]) for c in CODES}


def window_pv(pk, day):
    return np.array([pk.pv[s] for s in window_starts(day)])


def solve_call(pk, thresholds):
    ref, sundays, by_day = reference_day(pk)
    pv = window_pv(pk, ref)
    D = window_demand(pk, ref)
    hs = heaters(pk)
    L = routes(hs)
    F = {k: {c: D[c] - pk.outlook[c] / 1000.0 * pv + L[k][c] / 1000.0 for c in CODES} for k in L}
    S = dict(ref=ref, sundays=sundays, by_day=by_day, pv=pv, D=D, F=F, L=L, hs=hs, T=thresholds)
    S["mins"] = {k: {c: float(F[k][c].min()) for c in CODES} for k in F}
    S["sets"] = {k: sorted(c for c in CODES if S["mins"][k][c] < thresholds[c]) for k in F}
    S["designated"] = S["sets"]["G"]
    g = F["G"]
    S["curtail"] = {c: thresholds[c] - S["mins"]["G"][c] for c in S["designated"]}
    S["headroom"] = {c: float(np.min((g[c] - thresholds[c]) / pv)) for c in CODES if c not in S["designated"]}
    S["argmin"] = {c: float(HRS[int(np.argmin(g[c]))]) for c in CODES}
    S["n_heaters"] = {c: sum(1 for x in hs if x["sub"] == c) for c in CODES}
    S["n_reached"] = {c: sum(1 for x in hs if x["sub"] == c and x["cls"] == "R") for c in CODES}
    S["soak_mw_at_min"] = {c: float(L["G"][c][int(np.argmin(g[c]))] / 1000.0) for c in CODES}
    return S


# ---------------------------------------------------------------------------------------------- the asks
def solve_asks(pk, der_rows, on=P.REF_DAY):
    """Rooftop solar inverter capacity in service behind each zone substation on a date, and the part an
    emergency backstop can curtail, rebuilt from the applications register and the metering register."""
    alive = {}
    loc = {}
    for m in pk.meters:
        a, b = to_date(m["installed"]), to_date(m["removed"])
        loc[m["nmi"]] = m["zone_substation"]
        if a <= on and (b is None or b > on):
            alive[m["nmi"]] = True
    by_id = {r["application_id"]: r for r in der_rows}
    root = {}

    def root_of(aid):
        seen = []
        while True:
            r = by_id[aid]
            if r["application_type"] == "New connection" or not r["prior_application"]:
                break
            seen.append(aid)
            aid = r["prior_application"]
        for s in seen:
            root[s] = aid
        return aid
    chains = {}
    for r in der_rows:
        if r["status"] != "Commissioned":
            continue
        cd = to_date(r["commissioned"])
        if cd > on:
            continue
        chains.setdefault(root_of(r["application_id"]), []).append(r)
    cap = {c: 0.0 for c in CODES}
    bs = {c: 0.0 for c in CODES}
    for rt, links in chains.items():
        nmi = by_id[rt]["nmi"]
        if not alive.get(nmi):
            continue
        links.sort(key=lambda r: (to_date(r["commissioned"]), r["application_id"]))
        last = links[-1]
        inst = min(to_date(r["commissioned"]) for r in links if r["inverter_serial"] == last["inverter_serial"])
        kw = float(last["inverter_kw"])
        cap[loc[nmi]] += kw
        if inst and inst >= P.BACKSTOP_FROM:
            bs[loc[nmi]] += kw
    return {c: v / 1000.0 for c, v in cap.items()}, {c: v / 1000.0 for c, v in bs.items()}
