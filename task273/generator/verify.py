"""task273 architecture 5: the independent verifier. Reads only the bundle (target/, golden/, figures.json), on its own
code path, with no generator state: the window, thresholds and backstop date come off the framework PDF, the signal
times off the notice PDF, and every graded figure is recomputed and compared."""
from __future__ import annotations

import csv
import json
import os
import re
import subprocess
from collections import defaultdict
from datetime import date, datetime, timedelta

HERE = os.path.dirname(os.path.abspath(__file__))


def _pdf(path):
    t = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True).stdout
    return t


def _flat(t):
    return re.sub(r"\s+", " ", t)


def _clock(s):
    m = re.match(r"(\d{1,2})\.(\d{2}) (am|pm|noon)", s)
    h, mi, ap = int(m.group(1)), int(m.group(2)), m.group(3)
    if ap == "pm" and h != 12:
        h += 12
    return h + mi / 60.0


def _first_sunday_october(y):
    d = date(y, 10, 1)
    while d.weekday() != 6:
        d += timedelta(days=1)
    return d


def run(target, golden):
    fails = []

    def ok(name, cond, detail=""):
        print("%s  verify: %-56s %s" % ("PASS" if cond else "FAIL", name, detail))
        if not cond:
            fails.append(name)
    F = json.load(open(os.path.join(HERE, "figures.json")))
    files = {f: os.path.join(target, f) for f in os.listdir(target)}
    pick = lambda pat: next(p for f, p in files.items() if re.search(pat, f))

    # -------------------------------------------------------------- documents
    fw = _flat(_pdf(pick(r"framework.*\.pdf")))
    ok("framework PDF read (positive control)", "minimum-load threshold" in fw and "ten half-hours" in fw, "%d chars" % len(fw))
    m = re.search(r"ten half-hours from (\d+\.\d+ am) to (\d+\.\d+ pm)", fw)
    w0, w1 = _clock(m.group(1)), _clock(m.group(2))
    thr = {c: float(v.replace("−", "-")) for c, v in re.findall(r"\b([A-Z]{3}) [A-Z][a-z]+ (-\d+\.\d) ", fw)}
    ok("framework thresholds table read", len(thr) == 16, "%d thresholds" % len(thr))
    bm = re.search(r"installed on or after (\d{1,2}) (\w+) (\d{4})", fw)
    backstop = datetime.strptime("%s %s %s" % bm.groups(), "%d %B %Y").date()
    nt = _flat(_pdf(pick(r"notice.*\.pdf")))
    ok("notice PDF read (positive control)", "switch-on signal" in nt, "%d chars" % len(nt))
    rel = {int(g): _clock(t) for g, t in re.findall(r"Group (\d) (\d{1,2}\.\d{2} (?:am|pm|noon))", nt)}
    off = _clock(re.search(r"switch-off signal to every group at (\d{1,2}\.\d{2} (?:am|pm))", nt).group(1))
    ok("notice signal times read", len(rel) == 5, " ".join("%d:%.1f" % kv for kv in sorted(rel.items())))
    dst = _first_sunday_october(2026)

    def to_local(utc_end):
        start = utc_end - timedelta(minutes=30)
        change = datetime(dst.year, dst.month, dst.day, 2, 0) - timedelta(hours=10)
        return start + timedelta(hours=11 if start >= change else 10)

    # -------------------------------------------------------------- SCADA and the reference day
    dem = {}
    names = {}
    for r in csv.DictReader(open(pick(r"zss_net_demand"))):
        dem[(r["point_name"], to_local(datetime.fromisoformat(r["interval_end_utc"])))] = float(r["value_mw"])
    tags = sorted({k[0] for k in dem})
    tot = defaultdict(float)
    for (t, ls), v in dem.items():
        if w0 <= ls.hour + ls.minute / 60 < w1:
            tot[ls] += v
    sun = defaultdict(lambda: 1e9)
    for ls, v in tot.items():
        if ls.date().weekday() == 6:
            sun[ls.date()] = min(sun[ls.date()], v)
    ref = min(sun, key=sun.get)
    ok("reference day", ref.isoformat() == F["reference_day"], "%s, %.1f MW" % (ref, sun[ref]))
    ok("reference day network minimum", round(sun[ref], 1) == F["ref_network_min_mw"], "%.1f" % sun[ref])
    starts = [datetime(ref.year, ref.month, ref.day) + timedelta(minutes=int(30 * i)) + timedelta(hours=w0)
              for i in range(int((w1 - w0) * 2))]
    js = json.load(open(pick(r"pv_monitoring.*\.json")))
    pv = {}
    for v in js["values"]:
        pv[datetime.fromisoformat(v["interval_end"]).replace(tzinfo=None) - timedelta(minutes=30)] = v["output_kw_per_kw"]
    pvw = [pv[s] for s in starts]

    from openpyxl import load_workbook
    wb = load_workbook(pick(r"outlook.*\.xlsx"), read_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "Zone substation")
    hd = list(rows[hi])
    add = {r[hd.index("Code")]: float(r[hd.index("Forecast new inverter capacity (kW)")]) for r in rows[hi + 1:] if r[0]}
    tag_of = {}
    for t in tags:
        c = t.split("_")[0]
        if c not in add:
            c = next(k for k in add if F["per_substation"][k]["name"].upper().startswith(c))
        tag_of[c] = t

    # -------------------------------------------------------------- heaters
    wb = load_workbook(pick(r"meter_models"), read_only=True)
    mrows = list(wb.worksheets[0].iter_rows(values_only=True))
    mh = list(mrows[0])
    models = {}
    for r in mrows[1:]:
        d = dict(zip(mh, r))
        regs = d["Registers"]
        clr = re.search(r"(\w+) controlled load", regs)
        models[d["Model code"]] = (clr.group(1) if clr else None, "receiver" in d["Description"])
    cur, loc, ser_model = {}, {}, {}
    alive = set()
    for r in csv.DictReader(open(pick(r"metering_installations"))):
        ser_model[r["meter_serial"]] = r["model_code"]
        loc[r["nmi"]] = r["zone_substation"]
        if not r["removed"]:
            cur[r["nmi"]] = r["model_code"]
        a = date.fromisoformat(r["installed"])
        b = date.fromisoformat(r["removed"]) if r["removed"] else None
        if a <= ref and (b is None or b > ref):
            alive.add(r["nmi"])
    en = defaultdict(lambda: [0.0, 0])
    for r in csv.DictReader(open(pick(r"cl_customer_reads"))):
        if r["register_id"] == models[ser_model[r["meter_serial"]]][0]:
            en[r["nmi"]][0] += float(r["kwh"])
            en[r["nmi"]][1] += int(r["days"])
    soak = {c: [0.0] * len(starts) for c in add}
    every = {c: [0.0] * len(starts) for c in add}
    n_reached = 0
    for r in csv.DictReader(open(pick(r"load_control_register"))):
        n = r["nmi"]
        kw = float(r["element_kw"])
        t0 = rel[int(r["switching_group"])]
        dur = min(en[n][0] / en[n][1] / kw, off - t0)
        reach = models[cur[n]][1]
        n_reached += reach
        for i, s in enumerate(starts):
            h = s.hour + s.minute / 60.0
            ov = max(0.0, min(t0 + dur, h + 0.5) - max(t0, h)) / 0.5
            every[loc[n]][i] += kw * ov
            if reach:
                soak[loc[n]][i] += kw * ov
    ok("heaters the signal reaches", n_reached == F["heaters"]["reached"], "%d" % n_reached)
    des, des_every = [], []
    ps = F["per_substation"]
    for c in sorted(add):
        f = [dem[(tag_of[c], s)] - add[c] / 1000 * p + soak[c][i] / 1000 for i, (s, p) in enumerate(zip(starts, pvw))]
        fe = [dem[(tag_of[c], s)] - add[c] / 1000 * p + every[c][i] / 1000 for i, (s, p) in enumerate(zip(starts, pvw))]
        mn = min(f)
        ok("forecast minimum %s" % c, round(mn, 1) == ps[c]["forecast_min_mw"] and thr[c] == ps[c]["threshold_mw"],
           "%.3f against %.1f" % (mn, thr[c]))
        if mn < thr[c]:
            des.append(c)
            ok("curtailment %s" % c, round(thr[c] - mn, 1) == ps[c]["curtail_mw"], "%.3f" % (thr[c] - mn))
        else:
            h = min((x - thr[c]) / p for x, p in zip(f, pvw))
            ok("further capacity %s" % c, round(h, 1) == ps[c]["headroom_mw"], "%.3f" % h)
        if min(fe) < thr[c]:
            des_every.append(c)
    ok("the call", sorted(des) == F["call"]["designated"], ", ".join(des))
    ok("the default path's set", sorted(des_every) == F["call"]["default_path_set"], ", ".join(des_every))

    # -------------------------------------------------------------- asks
    wb = load_workbook(pick(r"der_applications"), read_only=True)
    it = wb.worksheets[0].iter_rows(values_only=True)
    dh = list(next(it))
    apps = {}
    for r in it:
        d = dict(zip(dh, list(r) + [None] * (len(dh) - len(r))))
        apps[d["application_id"]] = d

    def dd(v):
        return v.date() if isinstance(v, datetime) else v

    def root(a):
        while apps[a]["application_type"] != "New connection" and apps[a]["prior_application"]:
            a = apps[a]["prior_application"]
        return a
    chains = defaultdict(list)
    for a, d in apps.items():
        if d["status"] == "Commissioned" and dd(d["commissioned"]) <= ref:
            chains[root(a)].append(d)
    cap, bs = defaultdict(float), defaultdict(float)
    for rt, ls in chains.items():
        nmi = apps[rt]["nmi"]
        if nmi not in alive:
            continue
        ls.sort(key=lambda d: (dd(d["commissioned"]), d["application_id"]))
        last = ls[-1]
        first_seen = min(dd(d["commissioned"]) for d in ls if d["inverter_serial"] == last["inverter_serial"])
        cap[loc[nmi]] += float(last["inverter_kw"])
        if first_seen >= backstop:
            bs[loc[nmi]] += float(last["inverter_kw"])
    for c in sorted(add):
        ok("solar in service and backstop %s" % c, round(cap[c] / 1000, 2) == ps[c]["in_service_mw"] and
           round(bs[c] / 1000, 2) == ps[c]["backstop_mw"], "%.3f, %.3f" % (cap[c] / 1000, bs[c] / 1000))

    # -------------------------------------------------------------- goldens read back
    gx = [os.path.join(golden, f) for f in os.listdir(golden) if f.endswith(".xlsx")][0]
    wb = load_workbook(gx, data_only=True)
    got = {r[0]: r for r in wb.worksheets[0].iter_rows(min_row=5, values_only=True) if r[0] in add}
    ok("golden workbook carries every figure", all(got[c][3] == ps[c]["forecast_min_mw"] and got[c][4] ==
                                                   ("Yes" if ps[c]["designated"] else "No") for c in add), "16 rows")
    from docx import Document
    gd = [os.path.join(golden, f) for f in os.listdir(golden) if f.endswith(".docx")][0]
    text = " ".join(p.text for p in Document(gd).paragraphs)
    ok("golden briefing names every designated substation (positive control on the text)",
       "Recommendation" in text and all(ps[c]["name"] in text for c in F["call"]["designated"]), "%d chars" % len(text))
    print("verify: %d failures" % len(fails))
    return not fails


if __name__ == "__main__":
    TASK = os.path.dirname(HERE)
    run(os.path.join(TASK, "target"), os.path.join(TASK, "golden"))
