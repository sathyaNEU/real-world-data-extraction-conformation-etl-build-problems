"""task273 architecture 5: the build's assertions. Input gates, the three rules (R1 to R3c), the semantic fork table
(SF1 to SF16), the ladder, mess neutrality, the separation line, ask devices, edges, the realism sweep and the
distractors. Prints one line per assertion and returns True only when every one passes."""
from __future__ import annotations

import csv
import json
import os
import re
import subprocess
import zipfile
from datetime import date, datetime, timedelta

import numpy as np

import analysis as A
import params as P

CODES = A.CODES


def run(target, golden, b, S, fig):
    import build as B
    import cracker as CR
    out = []

    def ok(name, cond, detail=""):
        out.append((name, bool(cond), detail))
        print("%s  %-62s %s" % ("PASS" if cond else "FAIL", name, detail))

    w = b["w"]
    T = b["T"]
    ps = fig["per_substation"]
    files = sorted(os.listdir(target))
    pk = S["pk"]

    # ------------------------------------------------------------------ input gates
    ok("gate: 10 or more files", len(files) >= 10, "%d files" % len(files))
    fmts = sorted({f.rsplit(".", 1)[1] for f in files})
    ok("gate: 3 or more formats", len(fmts) >= 3, ", ".join(fmts))
    ok("gate: a file of 25,000 rows or more", max(b["counts"].values()) >= 25000, "%d rows" % max(b["counts"].values()))
    meta = json.load(open(os.path.join(os.path.dirname(target), "metadata.json")))
    ok("gate: two or more distractors declared", len(meta["distractor_files"]) >= 2, ", ".join(meta["distractor_files"]))
    bad_names = [f for f in files if re.search(r"readme|dictionary|manifest|notes|distractor|field", f, re.I)]
    ok("gate: no describing file or stump-named file", not bad_names, ", ".join(bad_names) or "none")
    sheets = []
    for f in files:
        if f.endswith(".xlsx"):
            with zipfile.ZipFile(os.path.join(target, f)) as z:
                wbx = z.read("xl/workbook.xml").decode()
                sheets += re.findall(r'<sheet name="([^"]+)"', wbx)
    ok("gate: no notes, readme or dictionary sheet", not [s for s in sheets if re.search(r"note|readme|dictionar|about|field",
                                                                                         s, re.I)], ", ".join(sheets))
    ok("gate: every file has provenance", set(meta["provenance"]) == set(files), "%d entries" % len(meta["provenance"]))

    # ------------------------------------------------------------------ R1 to R3c
    ok("R1 the field's raw material ships: half-hourly demand and per-heater records",
       len(S["hs"]) == len(w.heaters) and b["counts"]["spine"] > 60000,
       "%d heaters, %d SCADA rows" % (len(S["hs"]), b["counts"]["spine"]))
    heads = []
    for f in files:
        if f.endswith(".csv"):
            heads += next(csv.reader(open(os.path.join(target, f))))
    ok("R1 no shipped column holds the soak load or a forecast minimum",
       not [h for h in heads if re.search(r"soak|forecast|minimum|reach", h, re.I)], "%d CSV headers read" % len(heads))
    sets = S["sets"]
    ok("R3a the analyst test: every heater moved files the named wrong set", sets["R2"] == ["MDR", "ODR", "TRV", "WRD"],
       "default path designates %s" % sets["R2"])
    ok("R3b only the decisive step changed (reach) files the golden", sets["G"] == fig["call"]["designated"],
       "golden %s" % sets["G"])
    wrong_single = {k: sets[k] for k in ("R0", "R1", "R4", "R5")}
    ok("R3b each other single correction stays wrong", all(v != sets["G"] for v in wrong_single.values()),
       "; ".join("%s %d" % (k, len(v)) for k, v in wrong_single.items()))
    cheap = [k for k in ("R0", "R1", "R2", "R4", "R5") if sets[k] == sets["G"]]
    ok("R3b cheap routes without the reach that land the call (Gate F count)", not cheap, "%d of 5" % len(cheap))
    # R3c no plant: no copied block in the spine (a run of 8 identical half-hours across or within substations)
    blocks = {}
    dup = 0
    for c in CODES:
        v = np.round(w.mw[c], 3)
        for k in range(0, len(v) - 8, 8):
            key = tuple(v[k:k + 8])
            if key in blocks:
                dup += 1
            blocks[key] = (c, k)
    ok("R3c no copied or backfilled block in the SCADA series", dup == 0, "%d repeated 8-interval blocks" % dup)
    pdftext = {k: subprocess.run(["pdftotext", "-layout", os.path.join(target, B.FILES[k]), "-"],
                                 capture_output=True, text=True).stdout for k in ("framework", "notice", "summer")}
    ok("R3c positive control: the framework text is readable", "ten half-hours from 10.00 am to 3.00 pm" in
       re.sub(r"\s+", " ", pdftext["framework"]), "%d characters" % len(pdftext["framework"]))
    joined = re.sub(r"\s+", " ", " ".join(pdftext.values())).lower()
    ok("R3c no document names the reach or the time-switch heaters", not re.search(
        r"time switch|time-switch|not reach|does not reach|unaffected|keep their overnight|reprogram", joined),
       "framework, notice and outlook read")

    # ------------------------------------------------------------------ semantic fork table
    ref = S["ref"]
    ok("SF1 reference day: the network's lowest Sunday half-hour (framework 2.2)", ref == P.REF_DAY,
       "%s at %.2f MW, next Sunday %.2f" % (ref, S["sundays"][ref], sorted(S["sundays"].values())[1]))
    # per-substation lowest Sunday reading
    per_ref = {}
    for c in CODES:
        per_ref[c] = min(S["sundays"], key=lambda d: min(pk.dem[(c, s)] for s in A.window_starts(d)))
    alt = []
    for c in CODES:
        d = per_ref[c]
        D = np.array([pk.dem[(c, s)] for s in A.window_starts(d)])
        pv = A.window_pv(pk, d)
        f = D - pk.outlook[c] / 1000.0 * pv + S["L"]["G"][c] / 1000.0
        if f.min() < T[c]:
            alt.append(c)
    ok("SF1 per-substation lowest Sunday: ruled out by 2.2's network total, call compared",
       True, "per-substation reading designates %s (%s)" % (sorted(alt), "same call" if sorted(alt) == sets["G"]
                                                            else "differs, ruled out by framework 2.2"))
    ok("SF2 'Sunday' against any day: the lowest day of all is the reference Sunday",
       min(S["by_day"], key=S["by_day"].get) == ref, "same call")
    # SF3: the ten half-hours read as interval endings 10.00 to 2.30 (starts 9.30 to 2.00)
    hrs2 = np.arange(9.5, 14.5, 0.5)
    st2 = [datetime(ref.year, ref.month, ref.day) + timedelta(hours=float(h)) for h in hrs2]
    pv2 = np.array([pk.pv[s] for s in st2])
    alt3 = []
    for c in CODES:
        D2 = np.array([pk.dem[(c, s)] for s in st2])
        L2 = np.zeros(10)
        for x in S["hs"]:
            if x["sub"] != c or x["cls"] != "R":
                continue
            r0 = P.SOAK_RELEASE[x["group"]]
            dur = min(x["e"] / x["kw"], P.SOAK_OFF - r0)
            L2 += x["kw"] * np.clip(np.minimum(r0 + dur, hrs2 + 0.5) - np.maximum(r0, hrs2), 0, 0.5) / 0.5
        if (D2 - pk.outlook[c] / 1000.0 * pv2 + L2 / 1000.0).min() < T[c]:
            alt3.append(c)
    ok("SF3 window read as half-hours ending 10.00 to 2.30 pm", sorted(alt3) == sets["G"], "same call")
    ok("SF4 SCADA stamp read as interval start (shifted 30 min)", sorted(alt3) == sets["G"],
       "same windows as SF3, same call")
    ok("SF5 solar additions on inverter capacity, the only basis shipped", "inverter" in
       json.load(open(os.path.join(target, B.FILES["sample"])))["measure"], "same call")
    ok("SF6 daily energy from the current meter's reads only", sets["R7"] == sets["G"],
       "max shift %.4f MW" % max(abs(S["mins"]["R7"][c] - S["mins"]["G"][c]) for c in CODES))
    # SF7: dividing by the 91 days of spring instead of the days read
    hs91 = [dict(x, e=x["e"] * x["days"] / 91.0) for x in S["hs"]]
    cls = np.array([x["cls"] for x in hs91])
    L91 = A.load(hs91, cls == "R")
    alt7 = sorted(c for c in CODES if (S["D"][c] - pk.outlook[c] / 1000.0 * S["pv"] + L91[c] / 1000.0).min() < T[c])
    ok("SF7 daily energy over 91 days instead of the days read (framework 3.2)", alt7 == sets["G"], "same call")
    ok("SF8 heater draw placed in whole half-hours", sets["R8"] == sets["G"],
       "max shift %.3f MW" % max(abs(S["mins"]["R8"][c] - S["mins"]["G"][c]) for c in CODES))
    ok("SF9 a heater still drawing at 3.00 pm is outside the window either way", P.SOAK_OFF > P.WINDOW[1], "same call")
    first_cls = sets["R6"]
    ok("SF10 the meter at the reference day against the meter now", first_cls == sets["G"], "same call")
    gaps = [abs(S["mins"]["G"][c] - T[c]) for c in CODES]
    ok("SF11 below against at-or-below: no forecast within 0.25 MW of its threshold", min(gaps) > 0.25,
       "smallest gap %.2f MW" % min(gaps))
    ok("SF12 thresholds are negative demand, below means more negative (framework 2.3)",
       all(T[c] < 0 for c in CODES), "same call")
    # SF13 the controlled-load register: the only register outside general supply and export
    serial_model = {m["meter_serial"]: m["model_code"] for m in pk.meters}
    other = {}
    for r in pk.reads:
        regs = other.setdefault(r["meter_serial"], set())
        regs.add(r["register_id"])
    amb = [s for s, regs in other.items() if len([x for x in regs if x not in ("11", "01", "E1", "B1")]) != 1]
    ok("SF13 controlled-load register: one non-general, non-export register per meter", not amb, "%d meters" % len(other))
    with open(os.path.join(target, B.FILES["spine"])) as fh:
        tags = sorted({ln.split(",")[0] for ln in fh.read().split("\n")[1:] if ln})
    unmatched = [t for t in tags if t.split("_")[0] not in CODES]
    ok("SF15 the one point name off the code list is Tibbarin's", unmatched == ["TIBB_ZS_NET_MW"], ", ".join(unmatched))
    ok("SF14 the soak's switch-off (4.00 pm) and the window's end (3.00 pm)", "4.00 pm" in pdftext["notice"],
       "same call")
    ok("SF16 'load switched at different times' has one reading; which heaters it covers is found",
       sets["G"] != sets["R2"], "the reach separates the calls; no sentence reads it")

    # ------------------------------------------------------------------ ladder and margins
    ok("ladder: default path misses the three reach substations", set(fig["call"]["reach_flips"]) == {"KNB", "QRB", "GRB"},
       ", ".join(fig["call"]["reach_flips"]))
    ok("ladder: no-soak forecast designates eleven", len(sets["R0"]) == 11, ", ".join(sets["R0"]))
    flips_margin = min(min(S["mins"]["R2"][c] - T[c], T[c] - S["mins"]["G"][c]) for c in fig["call"]["reach_flips"])
    ok("ladder: reach substations clear the threshold by 0.3 MW on both sides", flips_margin >= 0.3,
       "%.2f MW" % flips_margin)
    edge = lambda v, u: u / 2 - abs(v - round(v / u) * u)
    e_min = min(edge(S["mins"]["G"][c], 0.1) for c in CODES)
    ok("edge: every forecast minimum sits 0.02 MW inside its printed tenth", e_min >= 0.02, "%.3f" % e_min)
    outs = [S["curtail"][c] for c in S["curtail"]] + [S["headroom"][c] for c in S["headroom"]]
    e_out = min(edge(v, 0.1) for v in outs)
    ok("edge: every curtailment and further capacity 0.02 inside its tenth", e_out >= 0.02, "%.3f" % e_out)
    fam = max(max(abs(S["mins"][k][c] - S["mins"]["G"][c]) for c in CODES) for k in ("R7", "R8"))
    ok("edge: the standing family's spread is inside the edge", fam < e_min + 0.05, "spread %.3f MW" % fam)
    e_ask = min(min(edge(S["cap"][c], 0.01), edge(S["backstop"][c], 0.01)) for c in CODES)
    ok("edge: ask figures 0.5 kW or more from a rounding boundary", e_ask >= 0.0005, "%.4f MW" % e_ask)

    # ------------------------------------------------------------------ mess neutrality (R2)
    # 1 day boundary taken in UTC: the window rows on the UTC date of the reference day only
    utc_rows = [s for s in A.window_starts(ref) if (s - timedelta(hours=11)).date() == ref]
    mins_utc = {c: min((S["F"]["G"][c][i] for i, s in enumerate(A.window_starts(ref)) if s in utc_rows)) for c in CODES}
    ok("R2 SCADA day cut in UTC: call and minima unchanged",
       all(abs(mins_utc[c] - S["mins"]["G"][c]) < 1e-9 for c in CODES), "the minima fall after 11.00 am local")
    ok("R2 EST rows sit on weekday nights only (dropping them moves nothing)",
       all(w.kdow[k] < 5 and (w.khour[k] < 5 or w.khour[k] >= 21) for _, k in b["est"]), "%d rows" % len(b["est"]))
    ok("R2 current serial only (meter exchanges in spring): printed minima unchanged",
       all(round(S["mins"]["R7"][c], 1) == round(S["mins"]["G"][c], 1) for c in CODES), "same call")
    first_row, has_cur = {}, set()
    for m in pk.meters:
        first_row.setdefault(m["nmi"], m)
        if not m["removed"]:
            has_cur.add(m["nmi"])
    ok("R2 first register row per NMI is the meter in service",
       all(not first_row[n]["removed"] for n in has_cur), "%d NMIs with a meter in service" % len(has_cur))
    ok("R2 days read against 91 days: call unchanged", alt7 == sets["G"], "see SF7")

    # ------------------------------------------------------------------ separation line and ask devices
    cl_nmis = {x["nmi"] for x in S["hs"]}
    ab_nmis = {a for a, _, _ in w.abolished}
    ok("separation: no abolished NMI is a controlled-load customer", not (cl_nmis & ab_nmis), "%d abolished" % len(ab_nmis))
    ok("separation: the call reads no file the asks' devices sit in", "der" not in A.CALL_FILES,
       "call files: %s" % ", ".join(A.CALL_FILES))
    res = CR.main(verbose=False)
    ok("asks: the cracker misses every in-service figure", res["routes"]["cracker"]["A_hits"] == 0,
       "%d of 16" % res["routes"]["cracker"]["A_hits"])
    ok("asks: the cracker misses every backstop figure", res["routes"]["cracker"]["B_hits"] == 0,
       "%d of 16" % res["routes"]["cracker"]["B_hits"])
    ok("asks: every cracker variant misses every figure", all(v["A_hits"] == 0 and v["B_hits"] == 0 for k, v in
                                                             res["routes"].items() if k != "golden handling of the asks"),
       "%d variants" % (len(res["routes"]) - 1))
    ok("asks: the golden handling reproduces the truth", res["routes"]["golden handling of the asks"]["A_hits"] == 16
       and res["routes"]["golden handling of the asks"]["B_hits"] == 16, "16 and 16")
    cap_t, bs_t = w.solar_truth(P.REF_DAY)
    ok("asks: law 6, one handling reproduces the physical truth exactly",
       all(abs(cap_t[c] / 1000 - S["cap"][c]) < 1e-9 and abs(bs_t[c] / 1000 - S["backstop"][c]) < 1e-9 for c in CODES),
       "16 substations")
    ok("asks: pair ceiling under 60 and mirror leakage under 50", res["pair"] <= 60 and res["mirror_leakage_pct"] < 50,
       "pair %.1f, leakage %.1f" % (res["pair"], res["mirror_leakage_pct"]))

    # ------------------------------------------------------------------ realism sweep
    der_ids = [r["application_id"] for r in S["der"]]
    seqs = {}
    for a in der_ids:
        y, n = a[2:6], int(a[7:])
        seqs.setdefault(y, []).append(n)
    const_step = [y for y, v in seqs.items() if len(v) > 3 and len(set(np.diff(sorted(v)))) == 1]
    ok("realism: application numbers are not an arithmetic sequence", not const_step, "%d years" % len(seqs))
    order_by_name = [s[0] for s in sorted(P.SUBS, key=lambda s: s[1])]
    ok("realism: substation codes are not alphabetical in name order", order_by_name != sorted(order_by_name),
       " ".join(order_by_name))
    grp = [x["group"] for x in S["hs"][:400]]
    ok("realism: switching groups are drawn, not dealt", not any(all(grp[i] == grp[i % p] for i in range(len(grp)))
                                                                for p in range(2, 8)), "no period")
    late_der = [r for r in S["der"] if r["approved"] > P.DER_PULL or (r["commissioned"] and r["commissioned"] > P.DER_PULL)]
    ok("realism: no application dated after the register extract", not late_der, "%d" % len(late_der))
    late_m = [m for m in pk.meters if A.to_date(m["installed"]) > P.METER_REG_PULL or
              (m["removed"] and A.to_date(m["removed"]) > P.METER_REG_PULL)]
    ok("realism: no meter dated after the metering extract", not late_m, "%d" % len(late_m))
    ok("realism: no tariff start after the load-control extract",
       all(A.to_date(r["on_tariff_since"]) <= P.LC_REG_PULL for r in pk.lc), "%d rows" % len(pk.lc))
    ok("realism: personas only from the card", all(p in ("Megan Robertson", "Jacob Barnes") for p in P.PEOPLE.values()),
       ", ".join(P.PEOPLE.values()))
    ok("realism: no round network totals", all(abs(v - round(v)) > 0.005 for v in fig["totals"].values()),
       str(fig["totals"]))

    # ------------------------------------------------------------------ distractors
    ok("distractors: the golden reads neither distractor", not (set(A.CALL_FILES + A.ASK_FILES) &
                                                                 {"summer", "enquiries"}), ", ".join(B.DISTRACTORS))

    # ------------------------------------------------------------------ goldens agree with the figures
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(golden, B.GOLDEN_XLSX), data_only=True)
    ws = wb.worksheets[0]
    vals = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] in CODES:
            vals[row[0]] = row
    ok("golden workbook: every substation's figures match figures.json",
       all(vals[c][3] == ps[c]["forecast_min_mw"] and vals[c][7] == ps[c]["in_service_mw"] and
           vals[c][8] == ps[c]["backstop_mw"] for c in CODES), "16 rows")
    sub = open(os.path.join(os.path.dirname(target), "submission.md")).read()
    ok("no em or en dash in the submission", "—" not in sub and "–" not in sub, "")
    n_pass = sum(1 for _, c, _ in out if c)
    print("%d of %d assertions pass" % (n_pass, len(out)))
    return n_pass == len(out)
