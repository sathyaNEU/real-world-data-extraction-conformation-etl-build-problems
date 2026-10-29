"""Build task273 (architecture 5): the evidence pack into <task>/target, the goldens into <task>/golden,
metadata.json, generator/figures.json and submission.md, then the checks and the independent verifier, then a
second build into a scratch folder inside the task (removed after) compared byte for byte, then target.zip.

    python3 task273/generator/build.py              everything
    python3 task273/generator/build.py --out DIR    the pack only, into DIR
    python3 task273/generator/build.py --golden     the pack and the goldens, no checks"""
import os
for _v in ("OPENBLAS_NUM_THREADS", "OMP_NUM_THREADS", "MKL_NUM_THREADS"):
    os.environ.setdefault(_v, "1")
import argparse      # noqa: E402
import calendar      # noqa: E402
import hashlib       # noqa: E402
import json          # noqa: E402
import math          # noqa: E402
import shutil        # noqa: E402
import sys           # noqa: E402
import warnings      # noqa: E402

import numpy as np   # noqa: E402

warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import analysis as A       # noqa: E402
import docs as DC          # noqa: E402
import params as P         # noqa: E402
import records as R        # noqa: E402
import world as Wd         # noqa: E402

TASK = os.path.dirname(HERE)
MTIME = calendar.timegm(P.ASSEMBLED.timetuple())
FILES = {"spine": "zss_net_demand_spring2026.csv",
         "sample": "rooftop_pv_monitoring_spring2026.json",
         "outlook": "solar_connections_outlook_2027.xlsx",
         "meters": "metering_installations_2026-12-11.csv",
         "models": "approved_meter_models.xlsx",
         "reads": "cl_customer_reads_spring2026.csv",
         "lc": "load_control_register_2026-12-10.csv",
         "der": "der_applications_2026-12-08.xlsx",
         "framework": "minimum_load_framework_v2.pdf",
         "notice": "network_notice_NN-2027-02.pdf",
         "summer": "summer_demand_outlook_2027-28.pdf",
         "enquiries": "voltage_enquiries_2026.csv"}
DISTRACTORS = ("summer", "enquiries")
GOLDEN_XLSX = "minimum_load_spring_2027.xlsx"
GOLDEN_DOCX = "minimum_load_briefing_2027.docx"
PDF_DATES = {"framework": (b"D:20260303101500+11'00'", b"D:20260303143000+11'00'"),
             "notice": (b"D:20270119084500+11'00'", b"D:20270119091200+11'00'"),
             "summer": (b"D:20261105161000+11'00'", b"D:20261106092500+11'00'")}
OOXML = {"outlook": (P.PEOPLE["ops"], P.PEOPLE["ops"], "2027-01-27T22:10:00Z", "2027-01-28T23:41:00Z",
                     "Rooftop solar connections outlook 2027", P.NETWORK),
         "models": ("Metering Services", "Metering Services", "2019-03-11T23:05:00Z", "2026-11-01T22:48:00Z",
                    "Approved meter models", P.NETWORK),
         "der": ("DER Connections", "DER Connections", "2026-12-07T21:02:00Z", "2026-12-07T21:02:00Z",
                 "Connection applications extract", P.NETWORK)}


def write_xlsx(path, sheets, title_rows=None, widths=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    for name, (head, rows) in sheets.items():
        ws = wb.create_sheet(name)
        for t in (title_rows or {}).get(name, []):
            ws.append([t])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=ws.max_row == 1)
        ws.append(head)
        hr = ws.max_row
        for c in ws[hr]:
            c.font = Font(bold=True)
        for row in rows:
            ws.append([None if v == "" else v for v in row])
        for i, h in enumerate(head):
            col = ws.cell(row=hr, column=i + 1).column_letter
            ws.column_dimensions[col].width = (widths or {}).get(i, max(10, min(60, 1.1 * max(
                len(str(h)), *(len(str(r[i])) for r in rows[:300])) + 1)) if rows else 12)
        ws.freeze_panes = "A%d" % (hr + 1)
    wb.save(path)


def write_der(path, head, rows):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Applications")
    hc = []
    for h in head:
        c = WriteOnlyCell(ws, value=h)
        c.font = Font(bold=True)
        hc.append(c)
    ws.append(hc)
    for i, w in enumerate((15, 12, 14, 16, 15, 15, 11, 10, 13, 12, 13)):
        ws.column_dimensions[chr(65 + i)].width = w
    for r in rows:
        out = []
        for v in r:
            if hasattr(v, "isoformat"):
                c = WriteOnlyCell(ws, value=v)
                c.number_format = "dd/mm/yyyy"
                out.append(c)
            else:
                out.append(None if v == "" else v)
        ws.append(out)
    wb.save(path)


# ---------------------------------------------------------------------------------------------- settle
def settle(w, S0):
    """Trim the network's connection forecast by a few tens of kW so each forecast minimum sits inside its printed
    tenth, then place each threshold on the tenth-of-a-megawatt grid by its ladder role."""
    D, pv, L = S0["D"], S0["pv"], S0["L"]
    Akw = dict(w.outlook_kw)
    rng = np.random.default_rng(w.seed + 71)

    def fc(k, c):
        return D[c] - Akw[c] / 1000.0 * pv + L[k][c] / 1000.0
    for _ in range(5):
        for c in w.codes:
            g = fc("G", c)
            m = float(g.min())
            target = round(m, 1) + float(rng.uniform(-0.012, 0.012))
            Akw[c] = int(round(Akw[c] + (m - target) * 1000.0 / pv[int(np.argmin(g))]))
    T = {}
    for c in w.codes:
        role = w.sub[c]["role"]
        g, r2, r0 = (float(fc(k, c).min()) for k in ("G", "R2", "R0"))
        lo, hi = {"F": (g + 0.30, r2 - 0.30), "D": (r2 + 0.30, r2 + 1.10), "C": (r0 + 0.30, g - 0.30),
                  "N": (r0 - 1.40, r0 - 0.30)}[role]
        cands = [round(x, 1) for x in np.arange(math.ceil(lo * 10) / 10.0, hi + 1e-9, 0.1)]
        assert cands, (c, role, lo, hi)
        mid = (lo + hi) / 2.0
        cands.sort(key=lambda t: abs(t - mid))
        pick = None
        for t in cands:
            if role in ("C", "N"):
                h = float(np.min((fc("G", c) - t) / pv))
                if 0.05 - abs(h - round(h, 1)) >= 0.02:
                    pick = t
                    break
            else:
                pick = t
                break
        assert pick is not None, (c, role)
        T[c] = float(pick)
    return Akw, T


# ---------------------------------------------------------------------------------------------- the pack
def build(out, seed=P.SEED, keep=False):
    if os.path.exists(out):
        shutil.rmtree(out)
    os.makedirs(out)
    p = lambda k: os.path.join(out, FILES[k])
    w = Wd.World(seed)
    Wd.build_reads(w)
    rc = R.Records(w)
    n = {}
    for key, fn in (("spine", rc.spine), ("meters", rc.meter_register), ("reads", rc.reads), ("lc", rc.load_control),
                    ("enquiries", rc.enquiries)):
        head, rows = fn()
        R.write_csv(p(key), head, rows)
        n[key] = len(rows)
    with open(p("sample"), "w", encoding="utf-8") as fh:
        fh.write(rc.sample_json())
        fh.write("\n")
    n["sample"] = w.K
    mh, mr = rc.model_rows()
    write_xlsx(p("models"), {"Approved models": (mh, mr)},
               title_rows={"Approved models": []}, widths={2: 92, 5: 52})
    n["models"] = len(mr)
    dh, dr = rc.der_rows()
    write_der(p("der"), dh, dr)
    n["der"] = len(dr)
    # the inputs the network's outlook and the Department's thresholds were settled against
    pk = A.Pack(out, FILES).load(partial=True)
    ref, _, _ = A.reference_day(pk)
    assert ref == P.REF_DAY, ref
    hs = A.heaters(pk)
    S0 = dict(D=A.window_demand(pk, ref), pv=A.window_pv(pk, ref), L=A.routes(hs))
    Akw, T = settle(w, S0)
    oh, orow = rc.outlook_rows(Akw)
    write_xlsx(p("outlook"), {"Outlook": (oh, orow)}, title_rows={"Outlook": [
        "Rooftop solar connections outlook for the 2027 minimum-load assessment",
        "Forecast to connect after the 2026 reference day and by 1 October 2027. Prepared by Network Planning, "
        "Tallowbark Networks, 29 January 2027."]})
    n["outlook"] = len(orow)
    DC.write_framework(p("framework"), T)
    DC.write_notice(p("notice"))
    DC.write_summer(p("summer"), w)
    return dict(w=w if keep else None, counts=n, A=Akw, T=T, est=sorted((c, k) for c, k in rc.est))


def finish(out):
    import postprocess as PP
    for k, (c, m) in PDF_DATES.items():
        PP.normalize_pdf(os.path.join(out, FILES[k]), c, m)
    for k, (cr, lb, c, m, title, co) in OOXML.items():
        PP.normalize_ooxml(os.path.join(out, FILES[k]), cr, lb, c, m, title=title, excel=True, company=co)
    for f in os.listdir(out):
        os.utime(os.path.join(out, f), (MTIME, MTIME))


def solve(target, T):
    pk = A.Pack(target, FILES).load()
    S = A.solve_call(pk, T)
    der = A.xlsx_rows(os.path.join(target, FILES["der"]))
    for r in der:
        for k in ("approved", "commissioned"):
            r[k] = A.to_date(r[k])
        r["prior_application"] = r["prior_application"] or ""
        r["inverter_serial"] = r["inverter_serial"] or ""
    cap, bs = A.solve_asks(pk, der)
    S["cap"], S["backstop"] = cap, bs
    S["der"] = der
    S["pk"] = pk
    return S


def r1(v):
    return float("%.1f" % (round(v * 10) / 10.0))


def r2(v):
    return float("%.2f" % (round(v * 100) / 100.0))


def figures(w, S, b):
    T = b["T"]
    des = S["designated"]
    order = sorted(P.SUBS, key=lambda s: s[1])
    names = {c: n for c, n, *_ in P.SUBS}
    des_names = [names[c] for c in sorted(des, key=lambda c: names[c])]
    fig = {
        "call": {"designated": sorted(des), "designated_names": des_names, "count": len(des),
                 "default_path_set": S["sets"]["R2"], "no_soak_set": S["sets"]["R0"]},
        "reference_day": S["ref"].isoformat(),
        "per_substation": {}, "rungs": {k: S["sets"][k] for k in S["sets"]},
        "thresholds": T, "outlook_kw": b["A"],
        "totals": {"cap_mw": r2(sum(S["cap"].values())), "backstop_mw": r2(sum(S["backstop"].values()))},
        "designated_totals": {"cap_mw": r2(sum(S["cap"][c] for c in des)),
                              "backstop_mw": r2(sum(S["backstop"][c] for c in des))},
        "ref_network_min_mw": r1(S["sundays"][S["ref"]]),
        "ref_day_of_month": S["ref"].day,
        "window_clock": ["10.00", "3.00"],
        "framework_clauses": ["2.1", "2.2", "2.3", "3.1", "3.2", "5.2"],
        "heaters": {"total": len(S["hs"]), "reached": sum(1 for x in S["hs"] if x["cls"] == "R"),
                    "time_switch": sum(1 for x in S["hs"] if x["cls"] == "T")},
    }
    fig["call"]["reach_flips"] = sorted((c for c in des if c not in S["sets"]["R2"]), key=lambda c: names[c])
    for c, nm, *_ in order:
        m = S["mins"]["G"][c]
        row = {"name": nm, "threshold_mw": T[c], "forecast_min_mw": r1(m), "forecast_min_exact": round(m, 4),
               "designated": c in des, "default_min_mw": r1(S["mins"]["R2"][c]),
               "in_service_mw": r2(S["cap"][c]), "backstop_mw": r2(S["backstop"][c]),
               "in_service_exact": round(S["cap"][c], 5), "backstop_exact": round(S["backstop"][c], 5),
               "soak_mw_at_min": round(S["soak_mw_at_min"][c], 3), "heaters": S["n_heaters"][c],
               "heaters_reached": S["n_reached"][c]}
        if c in des:
            row["curtail_mw"] = r1(S["curtail"][c])
            row["curtail_exact"] = round(S["curtail"][c], 4)
        else:
            row["headroom_mw"] = r1(S["headroom"][c])
            row["headroom_exact"] = round(S["headroom"][c], 4)
        fig["per_substation"][c] = row
    return fig


def write_goldens(golden, S, fig):
    import deliver as DV
    if os.path.exists(golden):
        shutil.rmtree(golden)
    os.makedirs(golden)
    png = DV.chart_png(fig)
    DV.write_workbook(os.path.join(golden, GOLDEN_XLSX), fig)
    DV.write_briefing(os.path.join(golden, GOLDEN_DOCX), fig, png)


def finish_golden(golden):
    import postprocess as PP
    PP.normalize_ooxml(os.path.join(golden, GOLDEN_XLSX), P.PEOPLE["director"], P.PEOPLE["director"],
                       "2027-03-08T23:40:00Z", "2027-03-17T05:22:00Z", title="Minimum-load assessment, spring 2027",
                       excel=True, company=P.DEPT)
    PP.normalize_ooxml(os.path.join(golden, GOLDEN_DOCX), P.PEOPLE["director"], P.PEOPLE["director"],
                       "2027-03-15T22:05:00Z", "2027-03-18T01:47:00Z", title="Minimum-load designations, spring 2027",
                       company=P.DEPT)
    for f in os.listdir(golden):
        os.utime(os.path.join(golden, f), (MTIME, MTIME))


SOURCES = {
    "spine": "Tallowbark Networks' SCADA historian export of each zone substation's net demand, half-hourly, UTC",
    "sample": "the Department's monitored sample of rooftop PV output per kW of inverter capacity, half-hourly",
    "outlook": "Tallowbark Networks' forecast of rooftop solar connections to 1 October 2027 by zone substation",
    "meters": "Tallowbark Networks' metering installations register, current and removed meters",
    "models": "Tallowbark Networks' approved meter models list",
    "reads": "Tallowbark Networks' spring 2026 consumption by meter register for its controlled-load customers",
    "lc": "Tallowbark Networks' load-control register of controlled-load installations",
    "der": "Tallowbark Networks' rooftop solar connection applications register",
    "framework": "the Department's minimum-load framework, version 2",
    "notice": "Tallowbark Networks' notice of new controlled-load switching times",
    "summer": "Tallowbark Networks' summer 2027-28 zone substation demand outlook",
    "enquiries": "Tallowbark Networks' customer voltage enquiries, January to November 2026",
}
PULLS = {"spine": P.SCADA_PULL, "sample": P.SAMPLE_PULL, "outlook": P.OUTLOOK_DATE, "meters": P.METER_REG_PULL,
         "models": P.MODEL_LIST_DATE, "reads": P.READS_PULL, "lc": P.LC_REG_PULL, "der": P.DER_PULL,
         "framework": P.FRAMEWORK_DATE, "notice": P.NOTICE_DATE, "summer": P.SUMMER_DATE,
         "enquiries": P.ENQUIRY_PULL}


def metadata(fig, counts):
    src = ("constructed by the task generator (seeded); the %s, %s, the %s, the places, people and records are "
           "fictional" % (P.DEPT, P.NETWORK, P.REGION))
    return {
        "task": "task273",
        "title": "%s: minimum-load designations for spring 2027" % P.DEPT,
        "domain": "Energy Systems & Utilities (load forecasting)",
        "objective": "Forecasting & Predictive Modeling",
        "distractor_files": ["target/" + FILES[k] for k in DISTRACTORS],
        "deliverables": [GOLDEN_XLSX, GOLDEN_DOCX],
        "provenance": {FILES[k]: {"source": "%s; %s" % (SOURCES[k], src), "date": PULLS[k].isoformat(),
                                  "licence": "CC0 1.0 (constructed for this task)"} for k in sorted(FILES)},
        "rows": {FILES[k]: v for k, v in sorted(counts.items())},
        "answer": fig["call"]["sentence"],
        "stated_inputs": sorted({"%.1f" % v for v in fig["thresholds"].values()}),
    }


def digest(folder):
    h = hashlib.sha256()
    for root, _, files in sorted(os.walk(folder)):
        for f in sorted(files):
            fp = os.path.join(root, f)
            h.update(os.path.relpath(fp, folder).encode())
            h.update(open(fp, "rb").read())
    return h.hexdigest()


def main_build(target, golden, seed=P.SEED):
    b = build(target, seed, keep=True)
    finish(target)
    S = solve(target, b["T"])
    fig = figures(b["w"], S, b)
    import writeup as WU
    fig["call"]["sentence"] = WU.call_sentence(fig)
    write_goldens(golden, S, fig)
    finish_golden(golden)
    with open(os.path.join(HERE, "figures.json"), "w") as f:
        json.dump(fig, f, indent=1, sort_keys=True)
        f.write("\n")
    WU.write(json.load(open(os.path.join(HERE, "figures.json"))))
    with open(os.path.join(TASK, "metadata.json"), "w") as f:
        json.dump(metadata(fig, b["counts"]), f, indent=1, sort_keys=True)
        f.write("\n")
    return b, S, fig


def make_zip(target, zpath):
    import zipfile
    files = sorted(os.listdir(target))
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for f in files:
            zi = zipfile.ZipInfo("target/" + f, date_time=P.ASSEMBLED.timetuple()[:6])
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            z.writestr(zi, open(os.path.join(target, f), "rb").read())


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out")
    ap.add_argument("--golden", action="store_true")
    a = ap.parse_args()
    if a.out:
        build(a.out)
        finish(a.out)
        print("pack written to", a.out)
        return
    target, golden = os.path.join(TASK, "target"), os.path.join(TASK, "golden")
    b, S, fig = main_build(target, golden)
    print("pack:", b["counts"])
    if a.golden:
        return
    import checks as CK
    import verify as VF
    ok = CK.run(target, golden, b, S, fig)
    vok = VF.run(target, golden)
    scratch = os.path.join(TASK, ".rebuild_check")
    b2 = build(os.path.join(scratch, "target"), keep=True)
    finish(os.path.join(scratch, "target"))
    same = digest(target) == digest(os.path.join(scratch, "target"))
    S2 = solve(os.path.join(scratch, "target"), b2["T"])
    fig2 = figures(b2["w"], S2, b2)
    import writeup as WU
    fig2["call"]["sentence"] = WU.call_sentence(fig2)
    write_goldens(os.path.join(scratch, "golden"), S2, fig2)
    finish_golden(os.path.join(scratch, "golden"))
    same_g = digest(golden) == digest(os.path.join(scratch, "golden"))
    shutil.rmtree(scratch)
    print("rebuild byte-identical: target %s, golden %s" % (same, same_g))
    if ok and vok and same and same_g:
        make_zip(target, os.path.join(TASK, "target.zip"))
        print("target.zip", hashlib.sha256(open(os.path.join(TASK, "target.zip"), "rb").read()).hexdigest()[:8])
    else:
        print("NOT zipped: checks %s, verify %s, rebuild %s/%s" % (ok, vok, same, same_g))
        sys.exit(1)


if __name__ == "__main__":
    main()
