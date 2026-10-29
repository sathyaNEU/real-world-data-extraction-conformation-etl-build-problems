"""task273 architecture 5: the golden deliverables. The assessment workbook and the one-page briefing to the
Minister, with its chart, all from the solved figures."""
from __future__ import annotations

import io
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.font_manager as fm   # noqa: E402
import matplotlib.pyplot as plt       # noqa: E402

import params as P                    # noqa: E402

FONT_DIR = next((d for d in ("/usr/share/fonts/truetype/liberation/", os.path.expanduser("~/Library/Fonts/"),
                              "/Library/Fonts/") if os.path.exists(d + "LiberationSans-Regular.ttf")),
                "/usr/share/fonts/truetype/liberation/")
for f in ("LiberationSans-Regular.ttf", "LiberationSans-Bold.ttf"):
    fm.fontManager.addfont(FONT_DIR + f)
plt.rcParams.update({"font.family": "Liberation Sans", "svg.hashsalt": "task273", "pdf.fonttype": 42})

NAVY, GREY, AMBER, RED = "#1f4e79", "#9aa5b1", "#c26b1a", "#7a1f1f"


def _order(F):
    ps = F["per_substation"]
    return sorted(ps, key=lambda c: ps[c]["forecast_min_mw"] - ps[c]["threshold_mw"])


def chart_png(F):
    ps = F["per_substation"]
    order = _order(F)[::-1]
    fig, ax = plt.subplots(figsize=(7.4, 5.6), dpi=160)
    for i, c in enumerate(order):
        r = ps[c]
        col = AMBER if r["designated"] else NAVY
        ax.barh(i, r["forecast_min_mw"], color=col, height=0.62)
        ax.plot([r["threshold_mw"]] * 2, [i - 0.42, i + 0.42], color=RED, lw=2.0)
        ax.text(r["forecast_min_mw"] - 0.2, i + 0.02, "%.1f" % r["forecast_min_mw"], va="center", ha="right",
                fontsize=7.0, color="#1c2330", bbox=dict(boxstyle="square,pad=0.1", fc="white", ec="none", alpha=0.8))
    ax.set_yticks(range(len(order)))
    ax.set_yticklabels(["%s (%s)" % (ps[c]["name"], c) for c in order], fontsize=7.8)
    lo = min(min(ps[c]["forecast_min_mw"], ps[c]["threshold_mw"]) for c in order)
    ax.set_xlim(lo - 2.6, 0)
    ax.axvline(0, color="#1c2330", lw=0.8)
    ax.set_xlabel("Forecast lowest demand, 10.00 am to 3.00 pm, 2027 reference Sunday (MW)\n"
                  "Negative demand is flow back into the subtransmission network", fontsize=7.6)
    ax.set_title("Seven zone substations fall below their minimum-load threshold once the solar soak starts",
                 fontsize=9.6, loc="left", color="#1c2330")
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    from matplotlib.lines import Line2D
    from matplotlib.patches import Patch
    ax.legend(handles=[Patch(color=AMBER, label="Designated: forecast below threshold"),
                       Patch(color=NAVY, label="Clears its threshold"),
                       Line2D([0], [0], color=RED, lw=2, label="Minimum-load threshold")],
              loc="upper center", bbox_to_anchor=(0.42, -0.13), ncol=3, fontsize=7.4, frameon=False)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def write_workbook(path, F):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    ps = F["per_substation"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Designations 2027"
    ws["A1"] = "Minimum-load assessment, spring 2027: Tallowbark Networks zone substations"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Forecast for the 2027 reference Sunday (the 25 October 2026 reference day carried forward under "
                "framework section 3) with the solar soak from 1 September 2027. MW; negative demand is flow back "
                "into the subtransmission network.")
    ws["A2"].font = Font(italic=True, size=9)
    head = ["Code", "Zone substation", "Threshold (MW)", "Forecast lowest demand, 10.00 am to 3.00 pm (MW)",
            "Designated", "Solar to curtail at the lowest half-hour (MW)",
            "Further inverter capacity before any half-hour falls below the threshold (MW)",
            "Rooftop solar inverter capacity in service, 25 Oct 2026 (MW)",
            "Of which the emergency backstop can curtail (MW)"]
    ws.append([])
    ws.append(head)
    for c in ws[4]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[4].height = 58
    for c in sorted(ps, key=lambda c: ps[c]["name"]):
        r = ps[c]
        ws.append([c, r["name"], r["threshold_mw"], r["forecast_min_mw"], "Yes" if r["designated"] else "No",
                   r.get("curtail_mw"), r.get("headroom_mw"), r["in_service_mw"], r["backstop_mw"]])
        row = ws.max_row
        for col, fmt in zip("CDFGHI", ("0.0", "0.0", "0.0", "0.0", "0.00", "0.00")):
            ws["%s%d" % (col, row)].number_format = fmt
        if r["designated"]:
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor="FBE5D0")
    ws.append([None, "Network", None, None, "%d of 16" % F["call"]["count"], None, None,
               F["totals"]["cap_mw"], F["totals"]["backstop_mw"]])
    tr = ws.max_row
    for cell in ws[tr]:
        cell.font = Font(bold=True)
    ws["H%d" % tr].number_format = ws["I%d" % tr].number_format = "0.00"
    for col, wdt in zip("ABCDEFGHI", (7, 15, 11, 17, 11, 15, 17, 18, 17)):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = "A4:I%d" % (tr - 1)
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws2 = wb.create_sheet("Basis")
    lines = [
        ("Reference day", "Sunday 25 October 2026: lowest network half-hour between 10.00 am and 3.00 pm of the "
                          "spring 2026 Sundays, %.1f MW (framework 2.2). SCADA converted from UTC to local time." %
         F["ref_network_min_mw"]),
        ("Solar to 1 Oct 2027", "Tallowbark outlook of 29 January 2027, priced at the 25 October output per kW of "
                                "inverter capacity in the monitored sample (framework 3.1a)."),
        ("Solar soak", "Network notice NN-2027-02: switch-on signal by group from 10.00 am to noon, off at 4.00 pm. "
                       "The signal reaches the %s heaters whose meter's controlled-load contactor is driven by an "
                       "integral audio-frequency receiver; the %s heaters on meters with an internal time switch "
                       "keep their overnight programme." % (format(F["heaters"]["reached"], ","),
                                                           format(F["heaters"]["time_switch"], ","))),
        ("Heater draw", "Each reached heater draws its element rating from its group's switch-on until it has "
                        "taken its spring 2026 daily energy (controlled-load register kWh over days read, all of "
                        "its meters; framework 3.2)."),
        ("Thresholds", "Framework section 4 (reviewed February 2026)."),
        ("Solar in service", "Applications chained through prior_application; the latest commissioned on or before "
                             "25 October 2026; NMIs with a meter in service that day; located by the metering "
                             "register. Backstop: the inverter serial in service first commissioned on or after "
                             "1 July 2025 (framework 5.2)."),
    ]
    ws2.append(["Item", "Basis"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for a, b in lines:
        ws2.append([a, b])
        ws2.cell(row=ws2.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws2.cell(row=ws2.max_row, column=1).alignment = Alignment(vertical="top")
    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 110
    wb.properties.creator = P.PEOPLE["director"]
    wb.properties.lastModifiedBy = P.PEOPLE["director"]
    wb.save(path)


def write_briefing(path, F, png):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    ps = F["per_substation"]
    des = sorted((c for c in ps if ps[c]["designated"]), key=lambda c: ps[c]["name"])
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(9.5)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin = Cm(1.6)
        sec.bottom_margin = Cm(1.4)
    for sec in doc.sections:
        for part, txt in ((sec.header, "OFFICIAL"), (sec.footer, "OFFICIAL    MB27/0412")):
            p = part.paragraphs[0]
            p.text = txt
            p.alignment = 1
            p.runs[0].font.size = Pt(8)
    h = doc.add_paragraph()
    r = h.add_run("Ministerial briefing MB27/0412  |  To: Minister for Energy  |  From: %s, %s  |  18 March 2027  |  "
                  "For decision by 31 March 2027" % (P.PEOPLE["director"], P.ROLES["director"]))
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x5a, 0x64, 0x70)
    t = doc.add_paragraph()
    r = t.add_run("Seven Tallowbark zone substations should be designated for spring 2027")
    r.bold = True
    r.font.size = Pt(13.5)
    r.font.color.rgb = RGBColor(0x1f, 0x4e, 0x79)
    names = ", ".join("%s (%.1f MW against %.1f MW)" % (ps[c]["name"], ps[c]["forecast_min_mw"],
                                                       ps[c]["threshold_mw"]) for c in des)
    p = doc.add_paragraph()
    p.add_run("Recommendation. ").bold = True
    p.add_run("Designate %s for spring 2027 and direct %s to have emergency backstop arrangements in place for "
              "them by 1 September 2027. On the reference Sunday their lowest demand between 10.00 am and 3.00 pm "
              "is forecast below threshold: %s. The other nine clear their thresholds."
              % (", ".join(ps[c]["name"] for c in des[:-1]) + " and " + ps[des[-1]]["name"], P.NETWORK, names))
    tab = doc.add_table(rows=1, cols=2)
    tab.style = "Table Grid"
    tab.rows[0].cells[0].text = ("That you designate the seven zone substations above under section 2 of the "
                                 "minimum-load framework for spring 2027.")
    tab.rows[0].cells[1].text = "Agreed / Not agreed"
    for cell in tab.rows[0].cells:
        cell.paragraphs[0].runs[0].font.size = Pt(8.5)
    p = doc.add_paragraph()
    p.add_run("The solar soak. ").bold = True
    p.add_run("From 1 September 2027 the network sends its controlled-load switch-on signals between 10.00 am and "
              "noon instead of overnight. Only heaters whose meter switches on that signal move: %s of the %s "
              "controlled-load heaters. The other %s are on older meters whose contactor runs from an internal time "
              "switch, so they keep heating overnight. They are concentrated behind %s, which is why those three "
              "stay below threshold even with the soak; counting every heater as moving would clear them."
              % (format(F["heaters"]["reached"], ","), format(F["heaters"]["total"], ","),
                 format(F["heaters"]["time_switch"], ","),
                 ", ".join(ps[c]["name"] for c in F["call"]["reach_flips"][:-1]) + " and " +
                 ps[F["call"]["reach_flips"][-1]]["name"]))
    doc.add_picture(io.BytesIO(png), width=Cm(15.6))
    cap = doc.add_paragraph("Figure 1. Forecast lowest demand on the reference Sunday against each zone substation's "
                            "minimum-load threshold.")
    cap.runs[0].font.size = Pt(7.5)
    p = doc.add_paragraph()
    p.add_run("Backstop reach. ").bold = True
    p.add_run("Behind the seven designated substations %.2f MW of rooftop solar was in service on 25 October 2026, "
              "of which %.2f MW is on inverters installed since 1 July 2025 that accept remote curtailment. The "
              "workbook gives each substation's curtailment need and every substation's figures."
              % (F["designated_totals"]["cap_mw"], F["designated_totals"]["backstop_mw"]))
    n0 = doc.add_paragraph("The reference Sunday is 25 October 2026, the lowest network half-hour of spring 2026 "
                           "between 10.00 am and 3.00 pm (framework 2.2), carried forward with the solar forecast to "
                           "connect by 1 October 2027 and the new signal times; nothing else is changed (framework "
                           "3.3). Negative demand is power flowing back into the subtransmission network.")
    n0.runs[0].font.size = Pt(7.5)
    n = doc.add_paragraph("Basis: minimum-load framework v2 (sections 2 to 5); Tallowbark notice NN-2027-02; "
                          "Tallowbark solar connections outlook, 29 January 2027; SCADA, metering, load-control and "
                          "connection records to December 2026. Contact: %s, %s." %
                          (P.PEOPLE["director"], P.ROLES["director"]))
    n.runs[0].font.size = Pt(7.5)
    doc.core_properties.author = P.PEOPLE["director"]
    doc.core_properties.title = "Minimum-load designations, spring 2027"
    doc.save(path)
