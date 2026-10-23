"""Readers of the cut pack, shared by the analysis and the cracker: every figure the build reports is computed from
the files in target/, never from the world."""
import csv
import datetime as dt
import os
import re
import subprocess
from zoneinfo import ZoneInfo

import common as K

PT = ZoneInfo(K.TZ)
UTC = ZoneInfo("UTC")
USE_V2 = {v: k for k, v in K.USE_V2.items()}
PRIVATE_V1 = {"01", "02", "03", "04"}
PRIVATE_V2 = {"PASSENGER CAR", "SPORT UTILITY", "PICKUP", "VAN"}
UTYPE = {"1": "MV", "2": "PK", "3": "PED", "4": "CYC", "MV IN TRANSPORT": "MV", "PARKED MV": "PK",
         "PEDESTRIAN": "PED", "PEDALCYCLIST": "CYC"}
SEV = {"1": "K", "2": "A", "3": "B", "4": "C", "5": "O", "K": "K", "A": "A", "B": "B", "C": "C", "O": "O"}
PTYPE = {"1": "DRIVER", "2": "PASSENGER", "3": "PED", "4": "CYC", "DRIVER": "DRIVER", "PASSENGER": "PASSENGER",
         "PEDESTRIAN": "PED", "BICYCLIST": "CYC"}
FSYS = {"Interstate": 1, "Other freeways and expressways": 2, "Other principal arterials": 3, "Minor arterials": 4,
        "Major collectors": 5, "Minor collectors": 6, "Local": 7}


def path(target, key):
    return os.path.join(target, K.FILES[key])


def rows(target, key):
    with open(path(target, key), encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def pdf_text(p):
    return subprocess.run(["pdftotext", "-layout", p, "-"], capture_output=True, text=True, check=True).stdout


# ------------------------------------------------------------------------------------------------ the road inventory
def inventory(target):
    import openpyxl
    wb = openpyxl.load_workbook(path(target, "inventory"), read_only=True)
    segs = {}
    for i, r in enumerate(wb["Segments"].iter_rows(values_only=True)):
        if i and r[0]:
            segs.setdefault(r[0], []).append((float(r[1]), float(r[2]), int(r[3]), r[4]))
    legacy = {}
    for i, r in enumerate(wb["Legacy route numbers"].iter_rows(values_only=True)):
        if i and r[0]:
            legacy[str(r[0])] = r[1]
    for v in segs.values():
        v.sort()
    return segs, legacy


def locate(segs, route, mp):
    for b, e, fc, ur in segs.get(route, ()):
        if b <= mp < e:
            return fc, ur
    return None


# ------------------------------------------------------------------------------------------------ crashes
def local_time(r):
    """The crash's local (Pacific) date and hour from the agency's own date, time and time zone; hour None when no
    time was recorded."""
    d = dt.date.fromisoformat(r["CRASH_DATE"])
    t = r["CRASH_TIME"].strip()
    if not t:
        return d, None, None
    if r["TIME_ZONE"] == "UTC":
        h, m = (int(x) for x in t.split(":"))
        x = dt.datetime(d.year, d.month, d.day, h, m, tzinfo=UTC).astimezone(PT)
        return x.date(), x.hour, x.minute
    if t.endswith(("AM", "PM")):
        hm, ap = t.split()
        h, m = (int(x) for x in hm.split(":"))
        h = h % 12 + (12 if ap == "PM" else 0)
        return d, h, m
    if ":" in t:
        h, m = (int(x) for x in t.split(":"))
        return d, h, m
    v = int(t)
    return d, v // 100, v % 100


def crashes(target, segs=None, legacy=None):
    if segs is None:
        segs, legacy = inventory(target)
    out = {}
    for r in rows(target, "crashes"):
        d, h, mi = local_time(r)
        v2 = r["FORM_VERSION"] == "PTCR-23"
        loc = None
        if v2 and r["F_SYSTEM"]:
            loc = (int(r["F_SYSTEM"]), r["URBAN_RURAL"])
        elif (not v2) and r["LOC_ROUTE"]:
            loc = locate(segs, legacy[r["LOC_ROUTE"]], float(r["MILEPOST"]))
        out[r["CRASH_ID"]] = dict(
            id=r["CRASH_ID"], v2=v2, agency=r["AGENCY"], rpt=r["AGENCY_REPORT_NO"], date=d, hour=h,
            dtype=None if d is None else ("we" if d.weekday() >= 5 else "wd"), year=d.year,
            fc=loc[0] if loc else None, ur=loc[1] if loc else None, group=K.GROUP_OF[loc] if loc else None,
            atf=int(r["AT_FAULT_UNIT"]) if r["AT_FAULT_UNIT"] else None, sev=SEV[r["CRASH_SEVERITY"]],
            fhe=r["FIRST_HARMFUL_EVENT"], units=[])
    return out


def operator_names(target):
    import openpyxl
    wb = openpyxl.load_workbook(path(target, "operators"), read_only=True)
    lab = {v: k for k, v in K.LABEL.items()}
    m = {}
    for i, r in enumerate(wb["Operators"].iter_rows(values_only=True)):
        if i and r[0]:
            for nm in str(r[6]).split(";"):
                m[nm.strip().lower()] = lab[r[2]]
    return m


def unit_use(u, v2, names):
    """The unit's use in the rule's terms: PV private drivers, a commission use, or another use."""
    if u["HIT_RUN_UNIT"] == "Y":
        return "UNK"
    code = u["VEH_USE"]
    body = u["BODY_TYPE"]
    if v2:
        use = USE_V2.get(code, "UNK")
        if use == "PV":
            return "PV" if body in PRIVATE_V2 else ("MC" if body == "MOTORCYCLE" else "PVO")
        return use
    if code == "0":
        return "PV" if body in PRIVATE_V1 else ("MC" if body == "10" else "PVO")
    if code in ("1", "2"):
        return names.get(u["CARRIER_NAME"].strip().lower(), "UNRES")
    return {"3": "OC", "4": "BUS", "5": "GV", "6": "EM", "9": "UNK"}.get(code, "UNK")


def load(target):
    """Crashes with their units, each unit's use and whether the officer recorded it at fault."""
    segs, legacy = inventory(target)
    cr = crashes(target, segs, legacy)
    names = operator_names(target)
    for u in rows(target, "units"):
        c = cr[u["CRASH_ID"]]
        no = int(u["UNIT_NO"])
        t = UTYPE[u["UNIT_TYPE"]]
        use = unit_use(u, c["v2"], names) if t in ("MV", "PK") else None
        if c["v2"]:
            fault = c["atf"] == no
        else:
            fault = u["CONTRIB_CODE"] not in ("", "00")
        c["units"].append(dict(no=no, type=t, use=use, fault=fault, fled=u["HIT_RUN_UNIT"] == "Y"))
    return cr


# ------------------------------------------------------------------------------------------------ travel
def tmr_vmt(target):
    """Table 1 of the travel monitoring report: annual VMT (millions) by functional system and area, 2021 to 2025."""
    txt = pdf_text(path(target, "tmr"))
    t1 = txt.split("Table 1.")[1].split("Table 2.")[0]
    out = {}
    pat = re.compile(r"^\s*(%s)\s+(Urban|Rural)\s+([\d,\.]+)\s+([\d,\.]+)\s+([\d,\.]+)\s+([\d,\.]+)\s+([\d,\.]+)\s*$"
                     % "|".join(re.escape(k) for k in FSYS))
    for line in t1.splitlines():
        m = pat.match(line)
        if m:
            key = (FSYS[m.group(1)], m.group(2)[0])
            out[key] = {y: float(m.group(3 + i).replace(",", "")) for i, y in enumerate(K.YEARS)}
    assert len(out) == len(K.FCLASS), len(out)
    return out


def stations(target):
    gnames = {i + 1: g for i, g in enumerate(K.GROUPS)}
    return {r["STATION_ID"]: gnames[int(r["FACTOR_GROUP"])] for r in rows(target, "stations")}


def atr_totals(target):
    """2025 volumes at each group's stations by (day type, hour), both directions, every day recorded."""
    st = stations(target)
    tot = {g: {(d, h): 0.0 for d in ("wd", "we") for h in range(24)} for g in K.GROUPS}
    with open(path(target, "atr"), encoding="utf-8", newline="") as f:
        for r in csv.DictReader(f):
            g = st[r["STATION_ID"]]
            d = dt.date.fromisoformat(r["COUNT_DATE"])
            dtp = "we" if d.weekday() >= 5 else "wd"
            for h in range(24):
                v = r["HR%02d" % (h + 1)]
                if v != "":
                    tot[g][(dtp, h)] += float(v)
    return tot


def atr_shares(target):
    tot = atr_totals(target)
    out = {}
    for g, cells in tot.items():
        s = sum(cells.values())
        out[g] = {k: v / s for k, v in cells.items()}
    return out


# ------------------------------------------------------------------------------------------------ the asks' files
def persons(target):
    out = {}
    for r in rows(target, "persons"):
        out.setdefault(r["CRASH_ID"], []).append(dict(
            unit=int(r["UNIT_NO"]), pno=int(r["PERSON_NO"]), ptype=PTYPE[r["PERSON_TYPE"]], seat=r["SEAT_POSITION"],
            age=int(r["AGE"]), sex=r["SEX"], sev=SEV[r["INJURY_SEVERITY"]]))
    return out


def fatalities(target):
    return rows(target, "deaths")


def costs(target):
    txt = pdf_text(path(target, "costs"))
    out = {}
    for code in "KABCO":
        m = re.search(r"^\s*%s\s+[A-Za-z ]+?\s+([\d,]+)\s*$" % code, txt, re.M)
        out[code] = int(m.group(1).replace(",", ""))
    return out
