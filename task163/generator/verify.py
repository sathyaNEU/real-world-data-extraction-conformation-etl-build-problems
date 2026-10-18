"""Independent verifier for task163: reads only the shipped bundle (target/) and the write-up's figures, on its own code
path (no generator module is imported), and recomputes the plan line, its rival routes and every graded ask figure.
It reads back every format the pack ships (JSON, CSV, XLSX, DOCX) and the filed sentences the call and the asks lean
on, each with a positive control on the same extractor.

    python3 task163/generator/verify.py [target]"""
import csv
import json
import os
import re
import sys
import zipfile
from collections import Counter, defaultdict

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
FAILS = []


def check(label, cond, detail=""):
    print("%s  %s%s" % ("ok  " if cond else "FAIL", label, (": " + detail) if detail else ""))
    if not cond:
        FAILS.append(label)


def month(y, m):
    return (y - 2021) * 12 + m - 1


def docx_paragraphs(path):
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    out = []
    for p in re.findall(r"<w:p[ >].*?</w:p>", xml, re.S):
        t = "".join(re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", p))
        if t:
            out.append(t.replace("&amp;", "&"))
    return out


def rows(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def main():
    target = sys.argv[1] if len(sys.argv) > 1 else os.path.join(TASK, "target")
    fig = json.load(open(os.path.join(HERE, "figures.json")))
    p = lambda f: os.path.join(target, f)

    # ---------------------------------------------------------------- formats read back
    api = json.load(open(p("bfs_state_api_2026-10-14.json")))
    check("JSON: the API pull reads, header first", api[0][0] == "cell_value" and len(api) > 8000, "%d rows" % len(api))
    std = docx_paragraphs(p("plan_2027_sizing_standard.docx"))
    notes = docx_paragraphs(p("first_time_employer_counting_notes.docx"))
    check("DOCX: positive control on the extractor (both titles read)",
          std[0].startswith("2027 operating plan") and notes[0].startswith("First-time employers"))
    est = [t for t in std if "will produce" in t]
    check("DOCX: the standard files the estimand once",
          len(est) == 1 and "2025 share in the plan workbook" in est[0] and "business formations within four quarters"
          in est[0] and "the twelve months to September 2026" in est[0] and "unadjusted counts" in est[0] and
          "nearest 100" in est[0])
    allstd = " ".join(std)
    check("DOCX: the standard names no method for the formations", not re.search(
        r"\b(rate|ratio|per application|latest|closed cohort|trend|grow\w*)\b", allstd.split("Sizing")[1].split("Timetable")[0]
        if "Sizing" in allstd else allstd))
    nt = " ".join(notes)
    for phrase in ("first payroll ever runs with us", "state it first employed in", "re-forms as a new entity",
                   "we paid wages for it, on any of its accounts, with a pay date in that quarter",
                   "calendar quarter of its first pay date", "gross wages we ran on our payroll"):
        check("DOCX: the counting notes file '%s'" % phrase, phrase in nt)
    wb = openpyxl.load_workbook(p("plan_2027_workbook.xlsx"), data_only=True)
    share = {}
    for r in wb["Share 2025"].iter_rows(values_only=True):
        if r and isinstance(r[0], str) and len(r[0]) == 2 and r[0].isupper() and r[0] != "Total":
            share[r[0]] = (int(r[2]), int(r[3]), float(r[4]))
    check("XLSX: the share tab reads (twelve states)", len(share) == 12)
    lw = openpyxl.load_workbook(p("state_filing_changes_log.xlsx"), data_only=True)["Log"]
    live = {}
    for r in lw.iter_rows(values_only=True):
        if r and r[3] == "Integrated EIN application":
            live[r[1]] = r[2]
    check("XLSX: the change log dates three integrated EIN filings", sorted(live) == ["MI", "OH", "PA"])

    # ---------------------------------------------------------------- the release
    fips = {"39": "OH", "42": "PA", "26": "MI", "18": "IN", "27": "MN", "55": "WI", "21": "KY", "54": "WV",
            "47": "TN", "29": "MO", "19": "IA", "51": "VA"}
    head = {h: i for i, h in enumerate(api[0])}
    BA, F4 = defaultdict(dict), defaultdict(dict)
    BAsa = defaultdict(dict)
    for r in api[1:]:
        st = fips[r[head["state"]]]
        y, m = map(int, r[head["time"]].split("-"))
        t = month(y, m)
        v = float(r[head["cell_value"]])
        if r[head["data_type_code"]] == "BA_BA":
            (BA if r[head["seasonally_adj"]] == "no" else BAsa)[st][t] = v
        elif r[head["data_type_code"]] == "BF_BF4Q" and r[head["seasonally_adj"]] == "no":
            F4[st][t] = v
    states = sorted(BA)
    last_closed = max(F4["OH"])
    check("the release closes formations at June 2025 applications", last_closed == month(2025, 6))
    ttm = range(month(2025, 10), month(2026, 9) + 1)
    for s in states:
        b = sum(F4[s][t] for t in range(month(2024, 7), month(2025, 6) + 1))
        if b != share[s][1]:
            check("share base ties for %s" % s, False)
    check("the share tab's formations base ties to the release", all(
        sum(F4[s][t] for t in range(month(2024, 7), month(2025, 6) + 1)) == share[s][1] for s in states))

    reform = sorted(live)
    others = [s for s in states if s not in reform]
    L = {s: month(live[s].year, live[s].month) for s in reform}

    def smm(s, t, hi):
        v = [BA[s][u] for u in range(0, hi + 1) if u % 12 == t % 12]
        return sum(v) / len(v)

    def ordinary(s, t):
        """The state's own level for the calendar month before go-live, moved with the nine states' applications
        against their level over the same months."""
        hi = L[s] - 1
        return smm(s, t, hi) * sum(BA[o][t] for o in others) / sum(smm(o, t, hi) for o in others)

    y = {}
    for s in states:
        hi = (L[s] - 1) if s in reform else last_closed
        y[s] = sum(F4[s][t] for t in range(0, hi + 1)) / sum(BA[s][t] for t in range(0, hi + 1))
    num = den = 0.0
    for s in reform:
        for t in range(L[s], last_closed + 1):
            g = ordinary(s, t)
            num += F4[s][t] - y[s] * g
            den += y[s] * (BA[s][t] - g)
    rel = num / den
    total = 0.0
    added = 0.0
    for s in states:
        apps = sum(BA[s][t] for t in ttm)
        if s in reform:
            g = sum(ordinary(s, t) for t in ttm)
            f = y[s] * (g + rel * (apps - g))
            added += apps - g
        else:
            f = y[s] * apps
        total += share[s][2] * f
    call = int(100 * int(total / 100.0 + 0.5))
    check("the plan line recomputes", call == fig["call"] and abs(total - fig["call_exact"]) < 0.06,
          "%.1f -> %d" % (total, call))
    check("the added applications' relative yield", abs(rel - fig["components"]["rel_yield"]) < 0.0005,
          "%.4f" % rel)
    check("the added applications in the twelve months, to the nearest 100,000",
          int(100000 * int(added / 100000 + 0.5)) == fig["components"]["added_ttm_about"] and
          int(round(added)) == fig["step_figures"]["added_ttm"], "%.0f" % added)
    edge = 50 - abs(total - call)
    check("the line sits a quarter bin or more from an edge", edge >= 25, "%.1f" % edge)
    # rivals
    r0 = sum(share[s][2] * sum(F4[s][t] for t in range(last_closed - 11, last_closed + 1)) /
             sum(BA[s][t] for t in range(last_closed - 11, last_closed + 1)) * sum(BA[s][t] for t in ttm)
             for s in states)
    check("the default path (latest closed rate carried) lands a full bin away", abs(r0 - total) >= 100,
          "%.1f" % r0)
    carry = sum(v[0] for v in share.values())
    check("carrying 2025's first-time employers lands a full bin away", abs(carry - total) >= 100, "%d" % carry)
    grow = carry * sum(BA[s][t] for s in states for t in ttm) / \
        sum(BA[s][t] for s in states for t in range(month(2025, 1), month(2025, 12) + 1))
    check("2025 grown with the twelve states' applications lands a full bin away", abs(grow - total) >= 100,
          "%.1f" % grow)
    sa_total = 0.0
    for s in states:
        apps_sa = sum(BAsa[s][t] for t in ttm)
        apps = sum(BA[s][t] for t in ttm)
        if s in reform:
            g = sum(ordinary(s, t) for t in ttm) * apps_sa / apps
            sa_total += share[s][2] * y[s] * (g + rel * (apps_sa - g))
        else:
            sa_total += share[s][2] * y[s] * apps_sa
    check("the seasonally adjusted twelve months give the same line", int(100 * int(sa_total / 100 + 0.5)) == call,
          "%.1f" % sa_total)

    # ---------------------------------------------------------------- the asks, on this code path
    acc = rows(p("onboarding_accounts_2022-2025.csv"))
    xw = rows(p("client_migration_crosswalk_2024-06-15.csv"))
    leg = rows(p("legacy_quarterly_wages_2022Q1-2024Q2.csv"))
    cur = rows(p("quarterly_wages_2024Q2-2025Q4.csv"))
    reg = rows(p("state_tax_registrations_2026-01-09.csv"))
    suta = rows(p("suta_successor_transfers_2022-2025.csv"))
    terr = {r["state"]: r["territory"] for r in rows(p("sales_territories_2025.csv"))}
    check("CSV: every company file reads", all((acc, xw, leg, cur, reg, suta, terr)))
    key = lambda e: re.sub(r"\D", "", e).rjust(9, "0")
    by_ein = defaultdict(list)
    for a in acc:
        by_ein[key(a["ein"])].append(a)
    regs = defaultdict(list)
    for r in reg:
        regs[r["account_id"]].append(r)
    succ_of = defaultdict(list)
    succ_eins = set()
    for r in suta:
        succ_of[key(r["predecessor_ein"])].append(key(r["successor_ein"]))
        succ_eins.add(key(r["successor_ein"]))
    acct_of = {r["legacy_client_no"]: r["account_id"] for r in xw}
    code_of = {r["account_id"]: r["legacy_source_code"] for r in xw}
    wq = Counter()
    for r in leg:
        if r["pay_group"] == "00":
            wq[(acct_of[r["client_no"]], int(r["tax_year"]), int(r["tax_qtr"]))] += int(r["gross_wages"])
    for r in cur:
        yy, qq = r["quarter"].split("-Q")
        wq[(r["account_id"], int(yy), int(qq))] += int(round(float(r["gross_wages"]) * 100))
    codes = {"ACCT": "accountant referral", "CPA": "accountant referral", "BANK": "bank partner", "BP": "bank partner",
             "WEB": "web sign-up", "OUT": "outbound", "TEL": "outbound", "MKT": "partner marketplace",
             "APX": "partner marketplace"}

    def paid(a, yy, qq):
        start = "%d-%02d-01" % (yy, 3 * qq - 2)
        return not (a["product"] == "Tax filing only" and a["product_since"] and a["product_since"] <= start)

    T = defaultdict(lambda: Counter())
    CH = Counter()
    for e, accts in by_ein.items():
        first = sorted(accts, key=lambda r: (r["first_pay_date"], r["account_id"]))[0]
        fp = first["first_pay_date"]
        if fp[:4] != "2024" or e in succ_eins:
            continue
        if any(r["liability_date"] < fp for a in accts for r in regs[a["account_id"]]):
            continue
        chain, todo, seen = [], [e], set()
        while todo:
            x = todo.pop()
            if x in seen:
                continue
            seen.add(x)
            chain += by_ein.get(x, [])
            todo += succ_of.get(x, [])
        st = min((r for a in accts for r in regs[a["account_id"]]),
                 key=lambda r: (r["liability_date"], r["filing_start"]))["state"]
        t = terr[st]
        T[t]["signed"] += 1
        if any(wq[(a["account_id"], 2025, 4)] > 0 and paid(a, 2025, 4) for a in chain):
            T[t]["still"] += 1
        yy, qq = int(fp[:4]), (int(fp[5:7]) - 1) // 3 + 1
        T[t]["wq"] += sum(wq[(a["account_id"], yy, qq)] for a in chain)
        T[t]["w24"] += sum(wq[(a["account_id"], 2024, q)] for a in chain for q in (1, 2, 3, 4) if paid(a, 2024, q))
        CH[codes[code_of[first["account_id"]]] if first["account_id"] in code_of else first["signup_channel"]] += 1
    for t, v in fig["territories"].items():
        check("ask: %s signed, still, first-quarter wages, 2024 wages" % t,
              T[t]["signed"] == v["signed"] and T[t]["still"] == v["still"] and
              "%.2f" % (T[t]["wq"] / 100) == v["first_quarter_wages"] and "%.2f" % (T[t]["w24"] / 100) == v["wages_2024"],
              "%d, %d, %.2f, %.2f" % (T[t]["signed"], T[t]["still"], T[t]["wq"] / 100, T[t]["w24"] / 100))
    check("ask: the channel table", all(CH[c] == n for c, n in fig["channels"].items()), str(dict(CH)))
    # the plain reads land elsewhere
    flagged = sum(1 for a in acc if a["first_time_employer"] == "Y" and a["first_pay_date"][:4] == "2024")
    check("the onboarding flag's 2024 count is not the golden's", flagged != sum(v["signed"] for v in
                                                                                 fig["territories"].values()),
          "%d flagged" % flagged)
    print("%d failures" % len(FAILS))
    sys.exit(1 if FAILS else 0)


if __name__ == "__main__":
    main()
