"""The goldens, from the shipped files only: first_time_employer_plan_2027.xlsx (the plan line, each state's working,
the added applications' yield, the 2024 first-time employers by territory and by signup channel),
first_time_employer_plan_2027_note.docx (one page: the figure and a chart of each territory's 2024 first-time employers
against those still with us), submission.md and generator/figures.json (every figure the write-up states, as it
states it)."""
from __future__ import annotations

import io
import json
import os

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import calc as CA
import common as K
import postprocess as PP

XLSX = "first_time_employer_plan_2027.xlsx"
DOCX = "first_time_employer_plan_2027_note.docx"
AUTHOR = K.PEOPLE["planning"]
NAMES = {"OH": "Ohio", "PA": "Pennsylvania", "MI": "Michigan", "IN": "Indiana", "MN": "Minnesota", "WI": "Wisconsin",
         "KY": "Kentucky", "WV": "West Virginia", "TN": "Tennessee", "MO": "Missouri", "IA": "Iowa", "VA": "Virginia"}
TERR_ORDER = ["Ohio", "Pennsylvania", "Michigan", "Indiana-Kentucky", "Upper Midwest", "South and Plains"]
CH_ORDER = ["accountant referral", "bank partner", "outbound", "partner marketplace", "web sign-up"]


def fmt(x, dp=0):
    v = K.round_half_up(x, dp)
    return format(v, ",.%df" % dp) if dp else format(int(v), ",")


def golive_text(X):
    names = {"OH": "Ohio", "PA": "Pennsylvania", "MI": "Michigan"}
    order = sorted(X.reform, key=lambda s: X.P.changes[s][0])
    return ", ".join("%s %s" % (names[s], K.us_date(X.P.changes[s][0]).replace(",", "")) for s in order)


def about_added(X):
    v = K.round_to(X.added_ttm, 100000)
    return "half a million" if v == 500000 else fmt(v)


def cents(c):
    return "$" + format(c / 100.0, ",.2f")


# ------------------------------------------------------------------------------------------------- the figures
class X_:
    pass


def compute(target):
    X = X_()
    P = CA.Pack(target)
    X.P = P
    X.G = CA.golden(P)
    X.total = X.G["total"]
    X.call = CA.round100(X.total)
    reform = sorted(P.changes)
    X.reform = reform
    X.plan_state = {s: P.share[s]["share"] * X.G["F"][s] for s in CA.STATES}
    X.added_ttm = sum(X.G["added_ttm"][s] for s in reform)
    X.ordinary_ttm = sum(X.G["ordinary_ttm"][s] for s in reform)
    X.apps_ttm = {s: CA.series(P, s, "BA_BA")[CA.TTM].sum() for s in CA.STATES}
    X.rel = X.G["rel"][reform[0]][1]
    X.y_reform = {s: X.G["yield"][s] for s in reform}
    X.F_total = sum(X.G["F"].values())
    X.reform_plan = sum(X.plan_state[s] for s in reform)
    # the added applications' yield, worked per state over the closed cohorts since go-live
    X.work = {}
    for s in reform:
        L = CA.launch_index(P.changes[s][0])
        G = CA.counterfactual(P, s, "CF1", [o for o in CA.STATES if o not in reform], L)
        BA, F4 = CA.series(P, s, "BA_BA"), CA.series(P, s, "BF_BF4Q")
        post = list(range(L, CA.LAST_CLOSED + 1))
        y = X.G["yield"][s]
        X.work[s] = {"months": len(post), "apps": BA[post].sum(), "ordinary": G[post].sum(),
                     "added": (BA[post] - G[post]).sum(), "formations": F4[post].sum(),
                     "from_ordinary": y * G[post].sum(), "yield": y,
                     "added_share_closed": (BA[post] - G[post]).sum() / BA[post].sum(),
                     "added_share_ttm": X.G["added_ttm"][s] / X.apps_ttm[s]}
    X.family = CA.family(P)
    B = CA.Book(target)
    X.T, X.CH, X.cohort = CA.asks(B)
    return X


# ------------------------------------------------------------------------------------------------- workbook
FILL = PatternFill("solid", fgColor="DCE4EC")
THIN = Side(style="thin", color="A9B4BE")
NAVY = "2F4F6B"


def _head(ws, row, cols, widths=None, height=30):
    for c, h in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = Font(name="Calibri", size=10, bold=True)
        cell.fill = FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = Border(bottom=THIN)
    ws.row_dimensions[row].height = height
    if widths:
        for c, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(c)].width = w


def _title(ws, title, lines, span):
    ws["A1"] = title
    ws["A1"].font = Font(name="Calibri", size=13, bold=True, color=NAVY)
    for i, t in enumerate(lines, start=2):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(name="Calibri", size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells("A%d:%s%d" % (i, span, i))
        ws.row_dimensions[i].height = 28


def _row(ws, r, vals, fmts, bold=False, top=False):
    for c, (v, f) in enumerate(zip(vals, fmts), start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(name="Calibri", size=10, bold=bold)
        if f:
            cell.number_format = f
        if top:
            cell.border = Border(top=THIN)


def write_xlsx(path, X):
    wb = Workbook()
    ws = wb.active
    ws.title = "Plan line"
    _title(ws, "2027 plan: first-time-employer line, twelve core states", [
        "Line for the plan: %s new first-time employers (to the nearest 100), from %s unrounded." %
        (fmt(X.call), fmt(X.total, 1)),
        "Each state: its 2025 share (plan workbook) times the formations within four quarters that the twelve months "
        "of applications to September 2026 will produce (BFS October 2026 release, not seasonally adjusted).",
        "Ohio, Pennsylvania and Michigan: the applications the integrated EIN filing added form employers at %.3f of "
        "the state's ordinary rate (see Added yield); the rest of their applications at the state's rate before "
        "go-live. The other nine: the state's rate over every closed cohort." % X.rel,
    ], "H")
    cols = ["State", "Applications, Oct 2025 to Sep 2026", "Ordinary applications", "Added by the integrated filing",
            "Formation rate (ordinary)", "Formations within 4 quarters, forecast", "2025 share", "Plan figure"]
    _head(ws, 6, cols, [16, 16, 14, 14, 12, 16, 10, 12], height=44)
    r = 7
    for s in CA.STATES:
        a = X.apps_ttm[s]
        if s in X.reform:
            g, i = X.G["ordinary_ttm"][s], X.G["added_ttm"][s]
        else:
            g, i = a, None
        _row(ws, r, [NAMES[s], a, g, i, X.G["yield"][s], X.G["F"][s], X.P.share[s]["share"], X.plan_state[s]],
             [None, "#,##0", "#,##0", "#,##0", "0.00%", "#,##0", "0.00%", "#,##0.0"])
        r += 1
    _row(ws, r, ["Twelve states", sum(X.apps_ttm.values()), None, X.added_ttm, None, X.F_total, None, X.total],
         [None, "#,##0", None, "#,##0", None, "#,##0", None, "#,##0.0"], bold=True, top=True)
    ws.freeze_panes = "A7"

    ws2 = wb.create_sheet("Added yield")
    _title(ws2, "The integrated filing's added applications: formation yield from the closed cohorts", [
        "Ordinary applications: the state's own average for each calendar month from January 2021 to go-live, moved "
        "month by month with the nine other states' applications against their average over the same months. Added: "
        "the rest, from the go-live month.",
        "Closed cohorts: applications from go-live to June 2025 (the last month the October release has formations "
        "within four quarters for). Relative yield, pooled: (formations less ordinary applications at the state's "
        "pre-go-live rate) over (added applications at that rate) = %.3f." % X.rel,
    ], "J")
    cols = ["State", "Go-live", "Closed months", "Applications", "Ordinary", "Added", "Added share",
            "Formations", "From ordinary at the pre-go-live rate", "Pre-go-live rate"]
    _head(ws2, 5, cols, [14, 12, 9, 13, 13, 12, 9, 12, 16, 10], height=44)
    r = 6
    for s in X.reform:
        w = X.work[s]
        _row(ws2, r, [NAMES[s], X.P.changes[s][0], w["months"], w["apps"], w["ordinary"], w["added"],
                      w["added_share_closed"], w["formations"], w["from_ordinary"], w["yield"]],
             [None, "yyyy-mm-dd", "0", "#,##0", "#,##0", "#,##0", "0.0%", "#,##0", "#,##0", "0.00%"])
        r += 1
    r += 1
    ws2.cell(row=r, column=1, value="Added share of applications, twelve months to September 2026").font = \
        Font(name="Calibri", size=10, bold=True)
    for s in X.reform:
        r += 1
        _row(ws2, r, [NAMES[s], None, None, None, None, None, X.work[s]["added_share_ttm"]],
             [None, None, None, None, None, None, "0.0%"])

    ws3 = wb.create_sheet("2024 first-time employers")
    _title(ws3, "First-time employers signed in 2024, by sales territory", [
        "Businesses whose first payroll ever ran with us in 2024, in the territory of the state they first employed "
        "in; still with us: we paid wages for the business with a pay date in October to December 2025; wages: gross "
        "wages we ran in the calendar quarter of the first pay date and in 2024 (January 2026 extract, legacy system "
        "through the crosswalk).",
    ], "E")
    _head(ws3, 4, ["Territory", "Signed in 2024", "Still with us, Q4 2025", "Gross wages, first quarter on payroll",
                   "Gross wages, 2024"], [20, 12, 14, 20, 18], height=32)
    r = 5
    for t in TERR_ORDER:
        v = X.T[t]
        _row(ws3, r, [t, v["signed"], v["still"], v["wages_cents"] / 100.0, v["wages_2024_cents"] / 100.0],
             [None, "#,##0", "#,##0", "$#,##0.00", "$#,##0.00"])
        r += 1
    tot = {k: sum(X.T[t][k] for t in TERR_ORDER) for k in ("signed", "still", "wages_cents", "wages_2024_cents")}
    _row(ws3, r, ["All territories", tot["signed"], tot["still"], tot["wages_cents"] / 100.0,
                  tot["wages_2024_cents"] / 100.0], [None, "#,##0", "#,##0", "$#,##0.00", "$#,##0.00"], bold=True,
         top=True)

    ws4 = wb.create_sheet("By channel")
    _title(ws4, "First-time employers signed in 2024, by signup channel", [
        "The channel of each business's first account; accounts opened on the legacy system by their legacy referral "
        "source code.",
    ], "B")
    _head(ws4, 4, ["Signup channel", "First-time employers, 2024"], [24, 14])
    r = 5
    for c in CH_ORDER:
        _row(ws4, r, [c, X.CH.get(c, 0)], [None, "#,##0"])
        r += 1
    _row(ws4, r, ["All channels", sum(X.CH.values())], [None, "#,##0"], bold=True, top=True)
    for w in (ws2, ws3, ws4):
        w.freeze_panes = "A%d" % (6 if w is ws2 else 5)
    ws5 = wb.create_sheet("Notes")
    ws5["A1"] = "Notes"
    ws5["A1"].font = Font(name="Calibri", size=13, bold=True, color=NAVY)
    notes = [
        "Prepared by Revenue Planning (FP&A) for the CFO's sign-off of the 2027 operating plan on 16 October 2026.",
        "Line: the 2027 sizing standard of 1 September 2026; shares from the plan workbook's Share 2025 tab "
        "(Sales Operations, refreshed 15 October 2026).",
        "Market: Census Bureau business formation statistics, October 2026 release (pulled 14 October 2026), not "
        "seasonally adjusted; formations within four quarters are published through June 2025 applications.",
        "Integrated EIN filings: go-live and closing dates from Market Research's log of state filing changes.",
        "2024 first-time employers: the year-end planning extract of 9 January 2026, counted as in Sales Operations' "
        "counting notes of 3 February 2026; wages in dollars and cents, legacy amounts converted from cents.",
        "Rounding: the line to the nearest 100 once the twelve states are added; counts whole; wages to the cent.",
    ]
    for i_, t in enumerate(notes, start=3):
        c = ws5.cell(row=i_, column=1, value=t)
        c.font = Font(name="Calibri", size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws5.row_dimensions[i_].height = 28
    ws5.column_dimensions["A"].width = 110
    for w in wb.worksheets:
        w.page_setup.orientation = "landscape"
        w.page_setup.fitToWidth = 1
        w.page_setup.fitToHeight = 0
        w.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)


# ------------------------------------------------------------------------------------------------- the note
def chart_png(X):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams["svg.hashsalt"] = "task163"
    fig, ax = plt.subplots(figsize=(6.4, 3.0), dpi=150)
    x = np.arange(len(TERR_ORDER))
    s = [X.T[t]["signed"] for t in TERR_ORDER]
    k = [X.T[t]["still"] for t in TERR_ORDER]
    b1 = ax.bar(x - 0.2, s, 0.4, color="#2F4F6B", label="Signed in 2024")
    b2 = ax.bar(x + 0.2, k, 0.4, color="#9DB3C6", label="Still with us, Q4 2025")
    for bars in (b1, b2):
        for bar in bars:
            ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + 6, "%d" % bar.get_height(), ha="center",
                    va="bottom", fontsize=7)
    ax.set_xticks(x)
    ax.set_xticklabels(TERR_ORDER, fontsize=7.5)
    ax.set_ylabel("Businesses", fontsize=8)
    ax.tick_params(axis="y", labelsize=7)
    share = [k[i] / s[i] for i in range(len(s))]
    title = ("About four in five of 2024's first-time employers were still with us at the end of 2025"
             if all(0.75 <= v < 0.85 for v in share) else "2024 first-time employers signed and still with us")
    ax.set_title(title, fontsize=9, loc="left")
    ax.legend(fontsize=7, frameon=False)
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def _page_field(par):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        r = par.add_run()
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        elif text.strip() == "PAGE":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            r._r.append(it)
        else:
            r.text = text


def write_docx(path, X, png):
    import docx
    from docx.shared import Cm, Pt, RGBColor
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    st.paragraph_format.space_after = Pt(5)
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(1.7), Cm(1.5)
    fp = sec.footer.paragraphs[0]
    fp.text = "Kilbourne Payroll FP&A, 2027 operating plan, page "
    _page_field(fp)
    grey = RGBColor(0x5D, 0x67, 0x6E)
    for r in fp.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = grey
    for label, value in (("To", "%s, CFO" % K.PEOPLE["cfo"]), ("From", "Revenue Planning, FP&A"),
                         ("Date", "16 October 2026"), ("Re", "First-time-employer line, 2027 operating plan")):
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        a = p.add_run(label + ": ")
        a.bold = True
        a.font.size = Pt(9)
        b = p.add_run(value)
        b.font.size = Pt(9)
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("%s first-time employers for 2027 in the twelve core states" % fmt(X.call))
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0x2F, 0x4F, 0x6B)
    d.add_paragraph(
        "We put %s new first-time employers in the 2027 plan's first-time-employer line. Ohio, Pennsylvania and "
        "Michigan now file the IRS employer number application from inside their online formation filings and have "
        "closed their older filing routes, so about %s of their %s applications in the twelve months to September "
        "2026 are ones the integrated filing added. Read against the nine other states before and after go-live, "
        "the added applications in the closed cohorts formed employers, taking the three states together, at under a "
        "third of the rate of the applications those states would have had anyway, and they are about a third of the "
        "three states' applications in the twelve months against about a fifth of the closed cohorts'. Every "
        "state's ordinary applications also rose about a tenth into the twelve months, from early 2026. Priced that "
        "way, the twelve "
        "months' applications will produce about %s employer formations within four quarters, and at the 2025 shares "
        "that is %s first-time employers, about four fifths of them in the three states." %
        (fmt(X.call), about_added(X), fmt(K.round_to(X.added_ttm + X.ordinary_ttm, 100000) / 1e6, 1) + " million",
         fmt(K.round_to(X.F_total, 5000)), fmt(X.call)))
    signed = sum(X.T[t]["signed"] for t in TERR_ORDER)
    still = sum(X.T[t]["still"] for t in TERR_ORDER)
    d.add_paragraph(
        "The 2024 cohort is the nearest guide to what the line is worth: %s first-time employers signed across the "
        "six territories, %s of them still with us in the fourth quarter of 2025. The workbook has each territory's "
        "first-quarter and 2024 wages and the channels the cohort came through." % (fmt(signed), fmt(still)))
    d.add_picture(io.BytesIO(png), width=Cm(16.0))
    p = d.add_paragraph()
    r = p.add_run("Source: year-end planning extract of 9 January 2026, counted as in Sales Operations' counting notes "
                  "of 3 February 2026 (a business once, in the territory of the state it first employed in).")
    r.font.size = Pt(8)
    r.font.color.rgb = grey
    p = d.add_paragraph()
    r = p.add_run("Shares are Sales Operations' 2025 shares in the plan workbook, refreshed from the BFS October 2026 "
                  "release on 15 October; the line is the twelve states' sum, rounded to the nearest 100 as the 2027 "
                  "sizing standard asks.")
    r.font.size = Pt(8)
    r.font.color.rgb = grey
    d.save(path)


# ------------------------------------------------------------------------------------------------- submission
def submission(X):
    T = X.T
    L = []
    L.append("## Tags\n")
    L.append("**Domain:** Economics & Econometrics (forecasting and forecast evaluation: an employer business "
             "formation forecast from business applications after three states' filing reforms changed who applies).")
    L.append("**Analytical objective:** Forecasting & Predictive Modeling.\n")
    L.append("## 1. Final Recommendation\n")
    L.append("**Put %s new first-time employers in the 2027 plan's first-time-employer line for the twelve core "
             "states.**" % fmt(X.call))
    L.append(
        "About %s of Ohio's, Pennsylvania's and Michigan's applications in the twelve months to September 2026 were "
        "added by their integrated EIN filings, and set against the nine other states before and after each go-live "
        "the added applications, pooled over the three states, form employers within four quarters at under a third "
        "of the ordinary rate; they are about a third of the three states' applications in the twelve months against "
        "about a fifth of the closed cohorts', and every state's ordinary applications rose about a tenth into the "
        "twelve months, from early 2026. Not each "
        "state's latest closed formation rate carried onto the twelve months, which prices the added applications at "
        "the share the closed cohorts held. Not the three states' own level before go-live as the counterfactual, "
        "which counts the early 2026 rise every state shared as added applications. Not the filing's effect on "
        "formations "
        "carried as a monthly level, which misses that the twelve months hold far more added applications than the "
        "closed cohorts. Not 2025's first-time employers grown with applications, which grows formations as fast as "
        "applications when the added applications that make up a third of the three states' twelve months form "
        "employers at under a third of the ordinary rate." % about_added(X))
    L.append("")
    L.append("## 2. Critical Components\n")
    L.append("1. About **%s** of the three states' applications in the twelve months to September 2026 were added by "
             "the integrated filing (to the nearest 100,000)." % about_added(X))
    L.append("2. Pooled over the three states, the added applications in the closed cohorts since go-live formed "
             "employers within four quarters at **under a third** of the rate of the applications the states would "
             "have had anyway.")
    L.append("3. The twelve states' applications in the twelve months will produce about **%s** employer formations "
             "within four quarters (to the nearest 5,000)." % fmt(K.round_to(X.F_total, 5000)))
    L.append("4. Ohio, Pennsylvania and Michigan carry about **four fifths** of the %s." % fmt(X.call))
    L.append("")
    L.append("## 3. Step-by-Step Solution\n")
    L.append(
        "1. Read plan_2027_sizing_standard.docx and the Share 2025 tab of plan_2027_workbook.xlsx: each state's line "
        "is its 2025 share times the business formations within four quarters that the twelve months of applications "
        "to September 2026 will produce, the twelve added and rounded to the nearest 100.")
    L.append(
        "2. Took the unadjusted BA_BA and BF_BF4Q rows of bfs_state_api_2026-10-14.json, whose formations close at "
        "June 2025 applications, and the integrated EIN filings' go-live dates in state_filing_changes_log.xlsx "
        "(%s)." % golive_text(X))
    L.append(
        "3. Set each of the three states' ordinary applications, month by month, at its own average for that "
        "calendar month from January 2021 to go-live times the nine other states' applications against their "
        "average for that calendar month over the same months, the rest from the go-live month being added: %s added "
        "of %s applications in the twelve months." %
        (fmt(X.added_ttm), fmt(X.added_ttm + X.ordinary_ttm)))
    L.append(
        "4. Divided the closed cohorts' formations since go-live, less the ordinary applications at each state's "
        "pre-go-live rate (Ohio %.2f, Pennsylvania %.2f, Michigan %.2f per cent), by the added applications at that "
        "rate, pooled over the three states: %.3f." % (100 * X.y_reform["OH"], 100 * X.y_reform["PA"],
                                                    100 * X.y_reform["MI"], X.rel))
    L.append(
        "5. Counted 2024 first-time employers by business as first_time_employer_counting_notes.docx counts them (an "
        "EIN's accounts, a client the legacy system re-opened under its corrected EIN, and its successors in "
        "suta_successor_transfers_2022-2025.csv as one; first paid wages in 2024 in the wage files, since a migrated "
        "account's first_pay_date is its legacy client-since date; no liability date before that quarter in "
        "state_tax_registrations_2026-01-09.csv), each in the territory in sales_territories_2025.csv of its earliest "
        "registration's state.")
    L.append(
        "6. Took wages from quarterly_wages_2024Q2-2025Q4.csv and from legacy_quarterly_wages_2022Q1-2024Q2.csv "
        "through client_migration_crosswalk_2024-06-15.csv at pay group 00 (the client's total), cents to dollars, "
        "leaving out quarters an account spent on Tax filing only, for the fourth quarter of 2025, each business's "
        "first quarter on payroll and 2024.")
    L.append(
        "7. Read each business's signup channel from its first account, an account opened on the legacy system by "
        "its legacy_source_code in the crosswalk.")
    L.append(
        "8. Priced the twelve months' ordinary applications in the three states at the state's pre-go-live rate and "
        "the added at "
        "%.3f of it, the nine states' applications at their rate over every closed cohort, and added each state's "
        "formations times its share: %s, which goes into the plan as %s." % (X.rel, fmt(X.total, 1), fmt(X.call)))
    L.append("")
    L.append("## 4. Deliverable Answers\n")
    L.append("### %s\n" % XLSX)
    L.append("1. The first-time-employer line for 2027, to the nearest 100: %s." % fmt(X.call))
    L.append("2. 2024 first-time employers by sales territory (signed in 2024; still with us in the fourth quarter "
             "of 2025; gross wages we ran for them in their first quarter on payroll; gross wages we ran for them in "
             "2024):")
    for t in TERR_ORDER:
        v = T[t]
        L.append("   - %s: %s signed, %s still with us, %s in the first quarter, %s in 2024" %
                 (t, fmt(v["signed"]), fmt(v["still"]), cents(v["wages_cents"]), cents(v["wages_2024_cents"])))
    L.append("3. 2024 first-time employers by signup channel:")
    for c in CH_ORDER:
        L.append("   - %s: %s" % (c, fmt(X.CH.get(c, 0))))
    L.append("")
    L.append("### %s\n" % DOCX)
    L.append("4. One page stating the line, %s new first-time employers for 2027, with the reason in two short "
             "paragraphs." % fmt(X.call))
    L.append("5. A bar chart under the text: two bars per sales territory, the 2024 first-time employers signed and "
             "those still with us in the fourth quarter of 2025, each labelled with its count as in item 2.")
    return "\n".join(L) + "\n"


def figures(X):
    T = X.T
    out = {
        "call": X.call,
        "call_exact": round(X.total, 1),
        "components": {"added_ttm_about": K.round_to(X.added_ttm, 100000),
                       "formations_ttm_about": K.round_to(X.F_total, 5000),
                       "reform_share_of_line": round(X.reform_plan / X.total, 3),
                       "rel_yield": round(X.rel, 3)},
        "step_figures": {"added_ttm": int(round(X.added_ttm)), "apps_ttm_reform": int(round(X.added_ttm +
                                                                                             X.ordinary_ttm)),
                         "unrounded_line": round(X.total, 1)},
        "steps": {"yield_OH_pct": round(100 * X.y_reform["OH"], 2), "yield_PA_pct": round(100 * X.y_reform["PA"], 2),
                  "yield_MI_pct": round(100 * X.y_reform["MI"], 2), "golden_exact": round(X.total, 1)},
        "territories": {t: {"signed": T[t]["signed"], "still": T[t]["still"],
                            "first_quarter_wages": "%.2f" % (T[t]["wages_cents"] / 100.0),
                            "wages_2024": "%.2f" % (T[t]["wages_2024_cents"] / 100.0)} for t in TERR_ORDER},
        "channels": {c: X.CH.get(c, 0) for c in CH_ORDER},
        "dates": [X.P.changes[s][0].isoformat() for s in X.reform],
        "roundings": [100, 5000, 100000],
    }
    return out


def main(target, golden, fig_path, sub_path):
    X = compute(target)
    os.makedirs(golden, exist_ok=True)
    xp, dp = os.path.join(golden, XLSX), os.path.join(golden, DOCX)
    write_xlsx(xp, X)
    write_docx(dp, X, chart_png(X))
    PP.normalize_ooxml(xp, AUTHOR, AUTHOR, "2026-10-15T21:05:00Z", "2026-10-16T13:22:00Z", excel=True)
    PP.normalize_docx(dp, AUTHOR, AUTHOR, "2026-10-16T12:40:00Z", "2026-10-16T13:31:00Z", revision=4, total_time=38,
                      seed="task163/note", title="2027 plan: the first-time-employer line")
    with open(sub_path, "w", encoding="utf-8") as fh:
        fh.write(submission(X))
    with open(fig_path, "w", encoding="utf-8") as fh:
        json.dump(figures(X), fh, indent=1, sort_keys=True)
        fh.write("\n")
    return X
