"""Writers for the data files in target/: the register workbook, the GL detail, the trial balances, the
capital approvals, the verification scans, the Bremer precedent load, the load specification and the
class mapping. Documents (PDF, EML, MD, TXT) are in docs.py."""
import csv
import datetime as dt
import os
import random
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import gl as GLM
import ledger as L
from params import (ACCOUNTS, GROUP_CLASSES_V2, LIFE_MONTHS, MAP_V2, MAP_V3, PRACTICE, SITES, TB_FINAL_RUN,
                    TB_PRELIM_RUN)

D = Decimal
Z = D("0")
BOLD = Font(bold=True)
HDR_FILL = PatternFill("solid", fgColor="DDEBF7")


def money(x):
    return "" if x == 0 else "%.2f" % x


def ddmmyyyy(d):
    return d.strftime("%d/%m/%Y") if d else ""


def widths(ws, ws_widths):
    for i, wdt in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt


# --------------------------------------------------------------------------- register workbook
def write_register(path, w):
    wb = Workbook()
    ws = wb.active
    ws.title = "Register"
    ws["A1"] = "YACAABA RADIOLOGY PTY LTD - FIXED ASSET REGISTER"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Maintained by Wright Accounting Partners. Copy for QIG migration, exported 2/10/2026"
    hdr = ["Asset No", "Tag", "Description", "Category", "Site", "Date Acquired", "Cost (ex GST)", "Supplier",
           "Invoice No", "Serial No", "Approval Ref", "Capitalised", "Status", "Date Disposed", "Notes"]
    ws.append([])
    ws.append(hdr)
    for c in ws[4]:
        c.font = BOLD
        c.fill = HDR_FILL
    rows = sorted(w.assets, key=lambda a: int(a["id"]))
    for a in rows:
        status = a["status"] if a["kind"] != "transfer_in" else "In use"
        ws.append([int(a["id"]), a["tag"] or None, a["description"], a["category"], SITES[a["site"]], a["acquired"],
                   float(a["cost"]), a["supplier"] or None, a["invoice"] or None, a["serial"] or None,
                   a["approval"] or None, a["capitalised"], status, a["disposed"], a["notes"] or None])
        r = ws.max_row
        ws.cell(r, 6).number_format = "dd/mm/yyyy"
        ws.cell(r, 7).number_format = "#,##0.00"
        if a["disposed"]:
            ws.cell(r, 14).number_format = "dd/mm/yyyy"
    widths(ws, [9, 9, 52, 24, 16, 13, 14, 32, 13, 13, 14, 11, 10, 13, 58])
    ws.freeze_panes = "A5"
    rt = wb.create_sheet("Rates")
    rt.append(["DEPRECIATION RATES (straight line)"])
    rt["A1"].font = BOLD
    rt.append([])
    rt.append(["Category", "Rate", "Life (years)", "Comment"])
    for c in rt[3]:
        c.font = BOLD
        c.fill = HDR_FILL
    rt.append(["Medical imaging equipment", 0.10, 10, None])
    rt.append(["Options & accessories", 0.20, 5, None])
    rt.append(["Ultrasound", 0.1667, 6, None])
    rt.append(["Computer equipment", 0.25, 4, "33.33% (3 yrs) to 30/06/2025"])
    rt.append(["Leasehold improvements", 0.0667, 15, None])
    rt.append(["Furniture & fittings", 0.10, 10, None])
    rt.append(["Motor vehicles", 0.1667, 6, None])
    for r in range(4, 11):
        rt.cell(r, 2).number_format = "0.00%"
    rt.append([])
    rt.append(["Full month's depreciation in the month of purchase. Items under $1,500 ex GST expensed (tracked on register for insurance, Capitalised = N)."])
    widths(rt, [28, 9, 12, 34])
    wb.save(path)


# --------------------------------------------------------------------------- GL detail
GL_HDR = ["Date", "Period", "Journal No", "Src", "Account No", "Account Name", "Job", "Memo", "Debit", "Credit"]


def write_gl(path, lines):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh)
        wr.writerow(GL_HDR)
        for r in lines:
            wr.writerow([ddmmyyyy(r["date"]), r["period"], r["journal"], r["src"], r["account"], r["account_name"],
                         r["job"], r["memo"], money(r["debit"]), money(r["credit"])])


def write_tb(path, prelim, final):
    wb = Workbook()
    first = True
    for title, rows, run in (("Run 03-10-2026 0914", prelim, "Printed 03/10/2026 9:14 AM"),
                             ("Run 09-10-2026 1602", final, "Printed 09/10/2026 4:02 PM")):
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = title
        ws["A1"] = "Yacaaba Radiology Pty Ltd"
        ws["A1"].font = BOLD
        ws["A2"] = "Trial Balance [Accrual] as at 30/09/2026"
        ws["A3"] = run
        ws.append([])
        ws.append(["Account No", "Account Name", "Debit", "Credit"])
        for c in ws[5]:
            c.font = BOLD
            c.fill = HDR_FILL
        dr = cr = Z
        for r in rows:
            ws.append([r["account"], r["name"], float(r["debit"]) if r["debit"] else None, float(r["credit"]) if r["credit"] else None])
            ws.cell(ws.max_row, 3).number_format = "#,##0.00"
            ws.cell(ws.max_row, 4).number_format = "#,##0.00"
            dr += r["debit"]
            cr += r["credit"]
        ws.append([None, "Total", float(dr), float(cr)])
        ws.cell(ws.max_row, 2).font = BOLD
        ws.cell(ws.max_row, 3).number_format = "#,##0.00"
        ws.cell(ws.max_row, 4).number_format = "#,##0.00"
        widths(ws, [12, 40, 16, 16])
    wb.save(path)


# --------------------------------------------------------------------------- capital approvals
def write_approvals(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh)
        wr.writerow(["Approval Ref", "Date Approved", "Site", "Request Type", "Description", "Existing Asset",
                     "Approved Amount (ex GST)", "Supplier", "Status", "Date Completed"])
        for r in rows:
            wr.writerow([r["ref"], ddmmyyyy(r["approved"]), SITES[r["site"]], r["type"], r["description"], r["existing"],
                         "%d" % r["amount"], r["supplier"], r["status"], ddmmyyyy(r["completed"])])


# --------------------------------------------------------------------------- verification scans
def write_scans(path, w):
    """The handheld export: each handheld numbers its own scans in the order it takes them, and the file
    lists them handheld by handheld. The time is whatever the handheld's clock read."""
    v = w.verification
    rng = random.Random(8110)
    base = {"HH-02": 41206, "HH-03": 52877, "HH-05": 70412, "HH-07": 18833}
    rows = []
    for a, site, t, tag, serial in v["scans"]:
        dev = rng.choice(v["devices"][site])
        usr = rng.choice(v["users"][site])
        base[dev] += rng.choice([1, 1, 2, 3])
        shown = v["clock"].get((a["id"] if a else None, t), t)
        rows.append((dev, base[dev], site, shown, tag, serial, usr))
    rows.sort(key=lambda r: (r[0], r[1]))
    with open(path, "w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh)
        wr.writerow(["scan_id", "scanner", "site_code", "scanned_at_utc", "tag", "serial", "user_id"])
        for dev, sid, site, t, tag, serial, usr in rows:
            wr.writerow([sid, dev, site, t.strftime("%Y-%m-%dT%H:%M:%SZ"), tag, serial, usr])


# --------------------------------------------------------------------------- precedent and group files
def write_bremer_load(path, rng):
    """Bremer's 1 July 2026 load file, the precedent the July migration produced. Bremer's ledger started
    depreciation the month after capitalisation and carried sub-numbered upgrades on their own five years."""
    fleet = [  # (class, life, description, cost range, sites, count)
        ("Imaging equipment", 120, "MRI 1.5T - Ingenia Ambition", (1.82e6, 2.05e6), ["IPS"], 1),
        ("Imaging equipment", 120, "CT 128-slice - SOMATOM go.Top", (1.05e6, 1.24e6), ["IPS", "SPF"], 2),
        ("Imaging equipment", 120, "General X-ray room - Multix Impact", (0.26e6, 0.36e6), ["IPS", "IPS", "SPF", "GDN", "GAT"], 5),
        ("Imaging equipment", 120, "Mammography - 3Dimensions", (0.52e6, 0.61e6), ["IPS"], 1),
        ("Imaging equipment", 120, "Fluoroscopy - Luminos Agile Max", (0.58e6, 0.66e6), ["IPS"], 1),
        ("Imaging equipment", 120, "OPG - Orthophos SL", (78e3, 96e3), ["SPF", "GDN"], 2),
        ("Imaging equipment", 120, "CT contrast injector - Stellant", (41e3, 48e3), ["IPS", "SPF"], 2),
        ("Clinical equipment", 72, "Ultrasound system - LOGIQ E10", (96e3, 205e3), ["IPS"] * 4 + ["SPF"] * 3 + ["GDN"] * 2 + ["GAT"], 10),
        ("Computer equipment", 48, "Workstation - Dell Precision 3660", (2.1e3, 3.3e3), ["IPS", "SPF", "GDN", "GAT"], 44),
        ("Computer equipment", 48, "Reporting monitor pair - Barco Nio", (9.6e3, 13.1e3), ["IPS", "SPF", "GDN"], 14),
        ("Leasehold improvements", 180, "Fit-out", (0.38e6, 2.05e6), ["IPS", "SPF", "GDN", "GAT"], 4),
        ("Furniture and fittings", 120, "Reception seating", (1.7e3, 4.4e3), ["IPS", "SPF", "GDN", "GAT"], 22),
        ("Furniture and fittings", 120, "Patient trolley", (3.9e3, 7.8e3), ["IPS", "SPF", "GDN", "GAT"], 18),
        ("Motor vehicles", 72, "Toyota HiLux", (44e3, 58e3), ["IPS", "SPF"], 2),
    ]
    rows = []
    no = 20500
    cut = L.mi(2026, 7)
    for cls, life, desc, (lo, hi), sites, k in fleet:
        for i in range(k):
            no += rng.choice([1, 2, 3])
            acq = dt.date(2014, 3, 1) + dt.timedelta(days=rng.randint(0, 4400))
            cost = D(int(rng.uniform(lo, hi))) + D(rng.randint(0, 99)) / 100
            st = L.mi(acq.year, acq.month) + 1
            s = L.sl(cost, life, st)
            acc = sum((x for m, x in s.items() if m < cut), Z)
            rem = max(0, max(s) - (cut - 1))
            site = sites[i % len(sites)]
            d_ = desc if cls != "Leasehold improvements" else "Fit-out %s" % {"IPS": "Ipswich", "SPF": "Springfield", "GDN": "Goodna", "GAT": "Gatton"}[site]
            rows.append(["BRM-%d" % no, "BRM-%d" % no, d_, site, cls, ddmmyyyy(acq), "%.2f" % cost, "%.2f" % acc,
                         "%.2f" % (cost - acc), "0.00", rem])
            if cls == "Imaging equipment" and cost > 250000 and acq.year < 2022 and rng.random() < 0.6:
                sub_acq = acq + dt.timedelta(days=rng.randint(500, 1500))
                if sub_acq < dt.date(2026, 4, 1):
                    sc = D(int(rng.uniform(70e3, 360e3))) + D(rng.randint(0, 99)) / 100
                    ss = L.sl(sc, 60, L.mi(sub_acq.year, sub_acq.month) + 1)
                    sa = sum((x for m, x in ss.items() if m < cut), Z)
                    srem = max(0, max(ss) - (cut - 1))
                    rows.append(["BRM-%d-1" % no, "BRM-%d-1" % no, "Upgrade - " + desc.split(" - ")[0], site, cls,
                                 ddmmyyyy(sub_acq), "%.2f" % sc, "%.2f" % sa, "%.2f" % (sc - sa), "0.00", srem])
    with open(path, "w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh)
        wr.writerow(["asset_number", "legacy_ref", "description", "site", "group_asset_class", "in_service_date", "cost",
                     "accumulated_depreciation", "net_book_value", "residual_value", "remaining_life_months"])
        for r in rows:
            wr.writerow(r)
    return rows


def write_load_spec(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Fields"
    ws["A1"] = "QIG fixed asset load file - field specification (FIN-SPEC-009 v1.2)"
    ws["A1"].font = BOLD
    ws.append([])
    ws.append(["Field", "Type", "Definition"])
    for c in ws[3]:
        c.font = BOLD
        c.fill = HDR_FILL
    rows = [
        ("asset_number", "text", "Asset number on the practice register (the legacy number)."),
        ("legacy_ref", "text", "Practice code and legacy number, e.g. YCB-1187."),
        ("description", "text", "Register description."),
        ("site", "code", "Site of record at cutover: CHA, MAI, CES or RAY for Yacaaba."),
        ("group_asset_class", "text", "Class the approved Group Asset Class Mapping gives the legacy account the asset is held in at the balances of record."),
        ("in_service_date", "date", "Date acquired as on the register (dd/mm/yyyy)."),
        ("cost", "AUD", "Legacy cost ex GST, to the cent."),
        ("accumulated_depreciation", "AUD", "Accumulated depreciation and impairment the practice ledger will hold for the asset at the close of the last legacy month, to the cent."),
        ("net_book_value", "AUD", "cost less accumulated_depreciation."),
        ("residual_value", "AUD", "Residual value the legacy ledger carried; 0.00 where none."),
        ("remaining_life_months", "integer", "Whole months from the cutover month over which the group ledger depreciates the asset; 0 where nothing remains to depreciate."),
    ]
    for r in rows:
        ws.append(list(r))
        ws.cell(ws.max_row, 3).alignment = Alignment(wrap_text=True)
    ws.append([])
    ws.append(["File: CSV, UTF-8, header row as above, one row per asset that loads. Bremer Radiology's 1 July 2026 file is the worked example."])
    widths(ws, [26, 9, 96])
    wb.save(path)


def write_mapping(path):
    wb = Workbook()
    practices = [("BRM", [("1-1810", "Medical Equipment at cost", "Imaging equipment"), ("1-1830", "Ultrasound Equipment at cost", "Clinical equipment"),
                          ("1-1850", "IT Equipment at cost", "Computer equipment"), ("1-1870", "Leasehold Improvements at cost", "Leasehold improvements"),
                          ("1-1890", "Furniture & Fittings at cost", "Furniture and fittings"), ("1-1910", "Motor Vehicles at cost", "Motor vehicles")]),
                 ("STZ", [("15100", "Imaging Plant & Equipment", "Imaging equipment"), ("15200", "Clinical Equipment", "Clinical equipment"),
                          ("15300", "Computers & Software", "Computer equipment"), ("15400", "Leasehold Improvements", "Leasehold improvements"),
                          ("15500", "Office Furniture", "Furniture and fittings"), ("15600", "Vehicles", "Motor vehicles")]),
                 ("YCB", [(k, GLM.NAMES[k], MAP_V2[k]) for k in sorted(MAP_V2)])]
    v3_change = {"Clinical equipment": "Imaging equipment", "Furniture and fittings": "Furniture, fittings and vehicles",
                 "Motor vehicles": "Furniture, fittings and vehicles"}
    for title, transform in (("v2", lambda c: c), ("v3", lambda c: v3_change.get(c, c))):
        ws = wb.active if title == "v2" else wb.create_sheet()
        ws.title = title
        ws["A1"] = "Group Asset Class Mapping %s" % title
        ws["A1"].font = BOLD
        ws.append([])
        ws.append(["Practice", "Legacy cost account", "Legacy account name", "Group asset class"])
        for c in ws[3]:
            c.font = BOLD
            c.fill = HDR_FILL
        for pr, accs in practices:
            for acc, name, cls in accs:
                ws.append([pr, acc, name, transform(cls)])
        ws.append([])
        ws.append(["Accumulated depreciation accounts follow their cost account."])
        widths(ws, [10, 20, 40, 34])
    lg = wb.create_sheet("Log")
    lg.append(["Version", "Issued", "Status", "Approved by", "Note"])
    for c in lg[1]:
        c.font = BOLD
        c.fill = HDR_FILL
    lg.append(["v1", "12/03/2026", "Superseded", "Group Financial Controller", "Bremer and Strzelecki accounts"])
    lg.append(["v2", "14/08/2026", "Approved", "Group Financial Controller", "Yacaaba accounts added for the January cutover"])
    lg.append(["v3", "30/09/2026", "Issued for review", "", "Proposed class consolidation for FY28 reporting; comments to R. Lawrence by 30/10"])
    widths(lg, [9, 12, 18, 26, 70])
    wb.save(path)
