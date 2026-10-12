"""The golden deliverables: Yacaaba's load file and internal audit's sign-off note, built from the
solution. verify_pack.py recomputes every figure in both from the shipped files."""
import csv
import io
import os
from decimal import Decimal

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from params import GROUP_CLASSES_V2, SITES
from solution import rnd

D = Decimal
SITE_ORDER = ["CHA", "MAI", "CES", "RAY"]
INK, INK2, SERIES = "#0b0b0b", "#52514e", "#2a78d6"


def write_load(path, sol):
    rows = sorted(sol["rows"], key=lambda r: int(r["asset"]))
    with open(path, "w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh)
        wr.writerow(["asset_number", "description", "site", "group_asset_class", "cost", "accumulated_depreciation",
                     "net_book_value", "remaining_life_months", "depreciation_jan_jun_2027"])
        for r in rows:
            wr.writerow([r["asset"], r["description"], r["site"], r["group_class"], "%.2f" % r["cost"], "%.2f" % r["acc_dec"],
                         "%.2f" % r["nbv"], r["remaining"], "%.2f" % r["h2"]])
    return len(rows)


def dollars(x):
    return "${:,}".format(rnd(x))


def chart_png(sol, path):
    vals = [rnd(sol["sites"][s]["h2"]) for s in SITE_ORDER]
    names = [SITES[s] for s in SITE_ORDER]
    fig, ax = plt.subplots(figsize=(6.6, 3.5), dpi=200)
    bars = ax.bar(names, vals, width=0.46, color=SERIES, zorder=3)
    top = max(vals)
    for b, v in zip(bars, vals):
        ax.text(b.get_x() + b.get_width() / 2, v + top * 0.015, "${:,}".format(v), ha="center", va="bottom",
                fontsize=8.5, color=INK)
    ax.set_title("January to June 2027 depreciation by site: ${:,} in total".format(rnd(sol["h2"])), fontsize=10, color=INK,
                 loc="left", pad=10)
    ax.set_ylabel("AUD", fontsize=8.5, color=INK2)
    ax.set_ylim(0, top * 1.14)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda x, p: "{:,.0f}".format(x)))
    ax.tick_params(colors=INK2, labelsize=8.5, length=0)
    ax.grid(axis="y", color="#e2e1dc", linewidth=0.6, zorder=0)
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.spines["bottom"].set_color("#b9b8b2")
    fig.tight_layout()
    fig.savefig(path, facecolor="white")
    plt.close(fig)


def _cell(c, text, bold=False, right=False, size=9, shade=None):
    c.text = ""
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if right:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if shade:
        tcpr = c._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:color"), "auto")
        sh.set(qn("w:fill"), shade)
        tcpr.append(sh)


def _table(doc, header, rows, widths, total=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(header):
        _cell(t.rows[0].cells[i], h, bold=True, right=i > 0, size=8.5, shade="DCE6F0")
    for r in rows + ([total] if total else []):
        cells = t.add_row().cells
        for i, v in enumerate(r):
            _cell(cells[i], v, bold=(r is total), right=i > 0, shade=("F2F2F2" if r is total else None))
    for row in t.rows:
        for i, wdt in enumerate(widths):
            row.cells[i].width = Cm(wdt)
    return t


def _source(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)


def _heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(11)


def _footer(doc):
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("Internal Audit  |  Yacaaba Radiology migration  |  Internal  |  Page ")
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    run = p.add_run()
    run.font.size = Pt(7.5)
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)


def write_note(path, sol, png):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(6)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
    _footer(doc)
    h = doc.add_paragraph()
    r = h.add_run("Yacaaba Radiology: load signed off, {} of depreciation January to June 2027".format(dollars(sol["h2"])))
    r.bold = True
    r.font.size = Pt(14)
    meta = doc.add_paragraph()
    mr = meta.add_run("To: Timothy Hill, Chief Financial Officer\nFrom: Cynthia Rice, Internal Audit Manager\n"
                      "Date: 23 October 2026\nFor: cutover go/no-go, 4 November 2026")
    mr.font.size = Pt(9.5)
    p = doc.add_paragraph()
    p.add_run("The group ledger will post {} of depreciation on Yacaaba's assets from January to June 2027".format(dollars(sol["h2"]))).bold = True
    p.add_run(", on the register as it loads at 1 January 2027 and with nothing bought or sold after cutover. "
              "I am signing the load off on that basis.")
    upg = [r for r in sol["rows"] if r["a"]["kind"] == "upgrade" and r["nbv"] == 0]
    months = sorted({sol["eng"].host_end(r["a"]) for r in upg})
    mname = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
    words = {1: "one", 2: "two", 3: "three", 4: "four", 5: "five", 6: "six", 7: "seven", 8: "eight", 9: "nine", 10: "ten", 11: "eleven", 12: "twelve"}
    doc.add_paragraph(
        "We rebuilt Yacaaba's schedules from the register and reproduce every depreciation posting in the Yacaaba ledger from July 2025 "
        "to September 2026, all six accounts and all fifteen months, to the cent, and the rebuild ties to the 9 October trial balance "
        "run, which is the run of record (FIN-STD-014 section 5). The rates tab carries no rate for equipment upgrades. Putting them on "
        "five years, the rate it gives options and accessories, replays the postings and the balances just as well, but it falls short "
        "of the 2024 depreciation the FY2025 report gives for imaging equipment in its comparative column. What reproduces that too is "
        "each upgrade spread over the months left on the system it was fitted to, and the capital approvals name that system for every "
        "upgrade. On that basis {} upgrades reach nil book value with their systems between {} and {} {} and come across at nil. "
        "Site is the site of record after the count (FIN-PRO-022 section 5). Raymond Terrace was counted on the paper sheet on "
        "7 October after its handheld failed, and the sheet is treated as the sighting record. {} assets that would be at nil book value "
        "and were not sighted by 5:00 pm on 9 October are written off on 31 December and are not in the load.".format(
            words.get(len(upg), str(len(upg))), mname[months[0] % 12], mname[months[-1] % 12], months[-1] // 12,
            words.get(len(sol["written_off"]), str(len(sol["written_off"]))).capitalize()))
    _heading(doc, "By site, as the assets load")
    rows = []
    tot = [D(0)] * 4
    for s in SITE_ORDER:
        v = sol["sites"][s]
        vals = [v["nbv_sep"], v["q4"], v["nbv_load"], v["h2"]]
        tot = [a + b for a, b in zip(tot, vals)]
        rows.append([SITES[s]] + [dollars(x) for x in vals])
    _table(doc, ["Site", "Book value 30 Sep 2026", "Oct to Dec 2026 depreciation", "Book value loaded 1 Jan 2027", "Jan to Jun 2027 charge"],
           rows, [3.4, 3.2, 3.4, 3.3, 3.2], total=["Total"] + [dollars(x) for x in tot])
    _source(doc, "Rebuilt schedules on Yacaaba's conventions; October to December is still posted by Yacaaba's ledger, January to June by the "
                 "group ledger under FIN-STD-014 section 6. Whole dollars.")
    _heading(doc, "Assets coming across at nil book value")
    nrows = []
    nc, ncost = 0, D(0)
    for s in SITE_ORDER:
        v = sol["sites"][s]
        nc += v["nil_count"]
        ncost += v["nil_cost"]
        nrows.append([SITES[s], str(v["nil_count"]), dollars(v["nil_cost"])])
    _table(doc, ["Site", "Assets", "Original cost"], nrows, [4.0, 2.2, 3.4], total=["Total", str(nc), dollars(ncost)])
    _source(doc, "Original cost as on the register, ex GST.")
    _heading(doc, "Book value loading by group asset class")
    crow = [[c, dollars(sol["classes"].get(c, D(0)))] for c in GROUP_CLASSES_V2]
    _table(doc, ["Group asset class", "Book value loaded"], crow, [6.0, 3.4],
           total=["Total", dollars(sum(sol["classes"].values(), D(0)))])
    _source(doc, "Group Asset Class Mapping v2 (approved 14 August 2026); v3 is still out for review. The Charlestown reception joinery sits "
                 "in leasehold improvements after the reclass in the 9 October run. The load file fa_load_2027-01-01.csv carries the "
                 "{} assets to the cent.".format(sol["load_rows"]))
    sig = doc.add_paragraph()
    sig.paragraph_format.space_before = Pt(10)
    sig.add_run("Cynthia Rice\nInternal Audit Manager, Quandong Imaging Group").font.size = Pt(10)
    doc.add_picture(png, width=Cm(15.5))
    doc.core_properties.author = "Cynthia Rice"
    doc.core_properties.title = "Yacaaba Radiology fixed asset migration sign-off"
    doc.save(path)


def build(sol, outdir, tmpdir):
    os.makedirs(outdir, exist_ok=True)
    n = write_load(os.path.join(outdir, "fa_load_2027-01-01.csv"), sol)
    png = os.path.join(tmpdir, "h2_by_site.png")
    chart_png(sol, png)
    write_note(os.path.join(outdir, "cutover_signoff.docx"), sol, png)
    return n
