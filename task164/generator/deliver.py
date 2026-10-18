"""The goldens from the shipped pack: ptr_2027_district_allocation.xlsx, summer_2026_transformer_failures.png,
generator/figures.json and submission.md. Every figure is computed by calc.py from <target> and carried here.

    python3 deliver.py [target golden figures.json submission.md]"""
from __future__ import annotations

import json
import os
import sys

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import calc as C  # noqa: E402
import postprocess as PP  # noqa: E402

TASK = os.path.dirname(HERE)
F_X = "ptr_2027_district_allocation.xlsx"
F_PNG = "summer_2026_transformer_failures.png"


def fmt_int(x):
    return "{:,}".format(int(x))


def fmt_usd(x):
    return "${:,}".format(int(round(x)))


def fmt_kva(x):
    return "{:,.1f}".format(x)


def figures(G, R):
    counts, alloc, q = G["counts"], G["alloc"], G["q"]
    A, B = G["A"], G["B"]
    F = {
        "call": {d: int(alloc[d]) for d in C.NAMES},
        "foothill_replacements": int(alloc["Foothill"]),
        "at_risk": {d: int(counts[d]) for d in C.NAMES},
        "at_risk_total": int(counts.sum()),
        "share": {d: round(float(q[d]), 3) for d in C.NAMES},
        "increment_default_kva": round(G["tab"]["tou"][("S", 1)], 1),
        "increment_ev_kva": round(G["tab"]["ev"], 1),
        "load_study_avg_default_kva": round(R["ls"]["avg_tou"], 2),
        "load_study_avg_ev_kva": round(R["ls"]["avg_ev"], 2),
        "default_path_foothill": int(R["R0"][1]["Foothill"]),
        "default_path_at_risk_total": int(R["R0"][0].sum()),
        "total_replacements": 400,
        "threshold_pct": 120,
        "threshold_ratio": 1.2,
        "ask_a": {d: {"replaced": int(A.loc[d, "replaced"]), "kva_added": round(float(A.loc[d, "kva_added"]), 1),
                      "cost_per_replacement": int(round(float(A.loc[d, "cost_per"])))} for d in C.PILOT},
        "ask_b": {d: {"customer_hours": int(round(float(B.loc[d, "customer_hours"]))),
                      "customers_over_8h": int(B.loc[d, "over_8h"])} for d in C.NAMES},
    }
    return F


def rungs(G):
    """The default path and the pooled increments, from the pack."""
    P, pn, Pt = G["P"], G["pn"], G["Pt"]
    ls = C.load_study_demand(P)
    f0 = C.constant_forecast(pn, ls["avg_tou"], ls["avg_ev"])
    c0 = C.at_risk(pn, f0)
    a0, _ = C.allocate(c0)
    return {"ls": ls, "R0": (c0, a0)}


def write_xlsx(path, G, F):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    head = PatternFill("solid", fgColor="1F4E79")
    hf = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    body = Font(name="Calibri", size=10)
    bold = Font(name="Calibri", bold=True, size=10)
    title = Font(name="Calibri", bold=True, size=13, color="1F4E79")
    sub = Font(name="Calibri", italic=True, size=9, color="555555")
    wrap = Alignment(wrap_text=True, vertical="top")

    def table(ws, top, cols, rows, widths, nfmt):
        for j, c in enumerate(cols):
            cell = ws.cell(row=top, column=j + 1, value=c)
            cell.fill, cell.font = head, hf
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if j else "left")
        ws.row_dimensions[top].height = 30
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                cell = ws.cell(row=top + 1 + i, column=j + 1, value=v)
                cell.font = body
                if nfmt[j]:
                    cell.number_format = nfmt[j]
        for j, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(j + 1)].width = w
        return top + 1 + len(rows)

    def note(ws, row, text, ncol, height):
        c = ws.cell(row=row, column=1, value=text)
        c.font, c.alignment = sub, wrap
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
        ws.row_dimensions[row].height = height

    A, inc = G["A"], G["tab"]
    ws = wb.active
    ws.title = "Allocation"
    ws["A1"] = "Foothill takes %d of the 400 replacements: %s transformers forecast above 120%% of nameplate for summer 2027" % (
        F["foothill_replacements"], fmt_int(F["at_risk_total"]))
    ws["A1"].font = title
    ws["A2"] = "2027 Proactive Service Transformer Replacement Program, D.26-11-019 Ordering Paragraphs 2 to 5. Record date 31 March 2027."
    ws["A2"].font = sub
    rows = [(d, F["at_risk"][d], F["share"][d], F["call"][d]) for d in C.NAMES]
    end = table(ws, 4, ["District", "Transformers at risk", "Proportional share", "Replacements"], rows,
                [17, 13, 13, 13], [None, "#,##0", "0.000", "#,##0"])
    ws.cell(row=end, column=1, value="System").font = bold
    for j, v, nf in ((2, F["at_risk_total"], "#,##0"), (3, 400.0, "0.000"), (4, sum(F["call"].values()), "#,##0")):
        c = ws.cell(row=end, column=j, value=v)
        c.font, c.number_format = bold, nf
    note(ws, end + 2, ("Forecast 2027 summer peak (OP 4): each transformer's recorded 2026 TLM peak, plus the effect of "
                       "each vehicle registered at its service points on 31 March 2027, less that of each vehicle "
                       "registered there on the day of its 2026 peak, each at its premise's schedule on that date. One "
                       "effect per schedule: the mean rise in recorded peak of the transformers that gained one vehicle on "
                       "that schedule and nothing else between the 2024, 2025 and 2026 summers, weather-adjusted by the "
                       "district's median ratio on transformers with no change."), 4, 76)
    note(ws, end + 3, ("At risk: forecast above 1.2 times the GIS kVA rating of the unit installed on the record date. "
                       "Shares to 400 by largest remainder (OP 2). Registrations dated 1 to 4 April 2027 in the 5 April "
                       "DMV feed are after the record date and not counted."), 4, 40)
    r0 = end + 5
    ws.cell(row=r0, column=1, value="Effect of one vehicle on a transformer's recorded summer peak, kVA").font = bold
    rows = [("Schedule R-EV2", inc["ev"]), ("Other residential schedules", inc["tou"][("S", 1)])]
    table(ws, r0 + 1, ["Vehicle at a service point on", "kVA per vehicle"], rows, [26, 13], [None, "0.00"])
    ws.column_dimensions["A"].width = 26
    ws.freeze_panes = "A5"
    ws.print_title_rows = "4:4"

    ws2 = wb.create_sheet("Pilot close-out")
    ws2["A1"] = "2025-26 pilot: %d transformers replaced in Railyard, Canal and Brookside" % int(A["replaced"].sum())
    ws2["A1"].font = title
    ws2["A2"] = "Decision 25-06-011 pilot, work completed 6 October 2025 to 24 April 2026. Charges to internal order 7100-PTR25 at 5 April 2027."
    ws2["A2"].font = sub
    rows = [(d, F["ask_a"][d]["replaced"], F["ask_a"][d]["kva_added"], round(float(A.loc[d, "total_cost"]), 2),
             F["ask_a"][d]["cost_per_replacement"]) for d in C.PILOT]
    table(ws2, 4, ["District", "Transformers replaced", "Nameplate kVA added", "Charges to pilot work orders ($)",
                   "Cost per transformer replaced ($)"], rows, [14, 13, 13, 17, 17],
          [None, "#,##0", "#,##0.0", "#,##0.00", "#,##0"])
    note(ws2, 9, ("Replaced: PTR changes in GIS whose installed serial differs from the removed serial, second units set "
                  "under another unit's order included. %d Railyard, %d Canal and %d Brookside units on the screen were "
                  "found larger on the pole than GIS recorded; their records were corrected under pilot orders and they "
                  "are not counted. Charges include cancelled and no-work orders, net of reversals." % tuple(
                      int(A.loc[d, "corrected"]) for d in C.PILOT)), 5, 64)
    ws2.freeze_panes = "A5"

    ws3 = wb.create_sheet("Summer 2026 failures")
    ws3["A1"] = "Transformer failures cost %s customer-hours between 1 June and 30 September 2026" % fmt_int(
        sum(v["customer_hours"] for v in F["ask_b"].values()))
    ws3["A1"].font = title
    ws3["A2"] = "OMS outage lines, cause EQ-XFMR, outage start in the window in local time (PDT, UTC less 7 hours)."
    ws3["A2"].font = sub
    rows = [(d, F["ask_b"][d]["customer_hours"], F["ask_b"][d]["customers_over_8h"]) for d in C.NAMES]
    end = table(ws3, 4, ["District", "Customer-hours", "Customers out more than 8 hours"], rows, [17, 14, 16],
                [None, "#,##0", "#,##0"])
    ws3.cell(row=end, column=1, value="System").font = bold
    for j, v in ((2, sum(v["customer_hours"] for v in F["ask_b"].values())), (3, sum(v["customers_over_8h"] for v in F["ask_b"].values()))):
        c = ws3.cell(row=end, column=j, value=v)
        c.font, c.number_format = bold, "#,##0"
    note(ws3, end + 2, ("Each outage line is counted at its own transformer's service points: the OMS customers_out field "
                        "repeats the event's total on every line of an event that took a tap or loop fuse. Hours run from "
                        "the outage start to the line's restoration; district from the GIS register."), 3, 64)
    ws3.freeze_panes = "A5"

    ws4 = wb.create_sheet("Notes")
    for line in ["Prepared for the D.26-11-019 Tier 2 advice letter due 30 April 2027.",
                 "Extracts pulled 5 April 2027: TLM season runs 2024 to 2026, GIS register and unit changes, service points,"
                 " rate schedule changes, DMV plug-in vehicle notifications, OMS outage lines, pilot work orders and cost lines.",
                 "Districts are operating districts in the GIS register on the record date.",
                 "Rounding: kVA to one decimal, cost to the nearest dollar, customer-hours to the nearest whole hour; "
                 "shares to three decimals for reference."]:
        ws4.append([line])
    ws4.column_dimensions["A"].width = 120
    for row in ws4.iter_rows():
        row[0].font = body
        row[0].alignment = wrap
    wb.save(path)
    PP.normalize_ooxml(path, "Katelyn Duarte", "Katelyn Duarte", "2027-04-12T16:05:00Z", "2027-04-14T22:31:00Z",
                       excel=True, company="Calloway Valley Power Company")


def write_png(path, F):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    names = list(C.NAMES)
    ch = [F["ask_b"][d]["customer_hours"] for d in names]
    o8 = [F["ask_b"][d]["customers_over_8h"] for d in names]
    order = sorted(range(len(names)), key=lambda i: ch[i])
    top = max(range(len(names)), key=lambda i: ch[i])
    fig, ax = plt.subplots(figsize=(9.4, 5.8), dpi=110)
    y = np.arange(len(names))
    cols = ["#1F4E79" if i == top else "#7FA7C9" for i in order]
    ax.barh(y, [ch[i] for i in order], color=cols, height=0.64)
    for yi, i in zip(y, order):
        ax.text(ch[i] + max(ch) * 0.012, yi, "{:,}   {} out over 8 h".format(ch[i], o8[i]), va="center", fontsize=8.6,
                color="#222222")
    ax.set_yticks(y)
    ax.set_yticklabels([names[i] for i in order], fontsize=9.4)
    ax.set_xlim(0, max(ch) * 1.38)
    ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: "{:,.0f}".format(v)))
    ax.set_xlabel("Customer-hours of interruption (each bar labelled with its customers out more than 8 hours)", fontsize=9)
    ax.set_title("%s's failed transformers cost the most customer-hours, summer 2026" % names[top], fontsize=11.5, loc="left",
                 color="#1F4E79")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.grid(axis="x", color="#e2e2e2", linewidth=0.6)
    ax.set_axisbelow(True)
    fig.text(0.01, 0.012, "Service transformer failures in service (OMS cause EQ-XFMR), 1 June to 30 September 2026; "
                          "customers from service points. Calloway Valley Power, PTR program.", fontsize=7.4, color="#555555")
    fig.tight_layout(rect=(0, 0.035, 1, 1))
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)


def submission(F):
    call = F["call"]
    lines = ["## Tags", "",
             "**Domain:** Energy Systems & Utilities (load forecasting).",
             "**Analytical objective:** Forecasting & Predictive Modeling.", "",
             "## 1. Final Recommendation", "",
             "**Allocate the 400 replacements as Foothill %d, Railyard %d, Canal %d, Mission %d, West Bluff %d, Crestview %d, "
             "Lakeside %d, Sunrise %d, Sandhill %d, Kingsley %d, Highland %d and Brookside %d.**" % tuple(call[d] for d in C.NAMES),
             "",
             "Foothill's %d of the %s transformers forecast above 120 percent of nameplate carry its %d replacements, each "
             "default-rate vehicle taken at the %.1f kVA one such vehicle adds to a transformer's recorded summer peak. Not the load "
             "study's 4 to 9 pm average of %.2f kVA a default-rate vehicle, which gives Foothill %d: it is the vehicle's mean "
             "draw over the window, not its effect on the transformer's highest quarter hour." % (
                 F["at_risk"]["Foothill"], fmt_int(F["at_risk_total"]), F["foothill_replacements"],
                 F["increment_default_kva"], F["load_study_avg_default_kva"], F["default_path_foothill"]),
             "", "## 2. Critical Components", "",
             "1. Transformers forecast at risk across the twelve districts: **%s**" % fmt_int(F["at_risk_total"]),
             "2. Foothill's transformers forecast at risk: **%d**" % F["at_risk"]["Foothill"],
             "3. Expected effect of one vehicle on a transformer's recorded summer peak, the mean weather-adjusted rise of "
             "the transformers that gained one vehicle on the schedule and nothing else between the 2024, 2025 and 2026 "
             "summers: **%.1f kVA** a default-rate vehicle, **%.1f kVA** an R-EV2 vehicle"
             % (F["increment_default_kva"], F["increment_ev_kva"]),
             "", "## 3. Step-by-Step Solution", "",
             "1. Placed each DMV notification at its matched_sp_id and that service point's transformer, a vehicle present on "
             "a date from reg_start_date up to its status_date, its premise on R-EV2 from the rate change effective_date.",
             "2. Counted each transformer's default-rate and R-EV2 vehicles on the dates of its 2024, 2025 and 2026 TLM "
             "peaks and on the record date of Ordering Paragraph 5.",
             "3. Took, per Ordering Paragraph 4, one effect per schedule: the mean rise in TLM peak of the transformers "
             "that gained one vehicle on that schedule and nothing else between summers, less the district's median ratio "
             "on transformers with no change: %.1f kVA a default-rate vehicle, %.1f kVA an R-EV2 one."
             % (F["increment_default_kva"], F["increment_ev_kva"]),
             "4. Forecast each transformer as its 2026 TLM peak plus those effects for its record-date vehicles less "
             "those for its peak-day vehicles, and counted those exceeding 1.2 times the register's kva_rating: "
             "%s, Foothill %d." % (fmt_int(F["at_risk_total"]), F["at_risk"]["Foothill"]),
             "5. Pilot: counted the PTR changes from 6 October 2025 to 24 April 2026 whose installed_serial differs from the "
             "removed_serial, their kVA change, and the cost lines on the district's pilot work orders, net of reversals, "
             "over those units.",
             "6. Failures: took the EQ-XFMR lines starting 1 June to 30 September local time (UTC less seven hours), each "
             "line in its transformer's register district with that transformer's service points as its customers and its "
             "hours to restored_utc.",
             "7. Split the 400 by largest remainder per Ordering Paragraph 2: Foothill %d." % F["foothill_replacements"],
             "", "## 4. Deliverable Answers", "", "### ptr_2027_district_allocation.xlsx", "",
             "1. Transformers forecast at risk and replacements by district:"]
    for d in C.NAMES:
        lines.append("   - %s: %s at risk, %d replacements" % (d, fmt_int(F["at_risk"][d]), call[d]))
    lines.append("2. Foothill's replacements: %d" % F["foothill_replacements"])
    lines.append("3. 2025-26 pilot by district:")
    for d in C.PILOT:
        a = F["ask_a"][d]
        lines.append("   - %s: %d transformers replaced, %s kVA added, %s per transformer replaced" % (
            d, a["replaced"], fmt_kva(a["kva_added"]), fmt_usd(a["cost_per_replacement"])))
    lines.append("4. Service transformer failures from 1 June to 30 September 2026 by district:")
    for d in C.NAMES:
        b = F["ask_b"][d]
        lines.append("   - %s: %s customer-hours, %s customers out more than eight hours" % (
            d, fmt_int(b["customer_hours"]), fmt_int(b["customers_over_8h"])))
    lines += ["", "### summer_2026_transformer_failures.png", "",
              "1. Horizontal bar chart of the customer-hours by district, each bar labelled with its customer-hours and "
              "its customers out more than eight hours.", ""]
    return "\n".join(lines)


def main(target=None, golden=None, fig_path=None, sub_path=None):
    target = target or os.path.join(TASK, "target")
    golden = golden or os.path.join(TASK, "golden")
    fig_path = fig_path or os.path.join(HERE, "figures.json")
    sub_path = sub_path or os.path.join(TASK, "submission.md")
    G = C.golden(target)
    R = rungs(G)
    F = figures(G, R)
    os.makedirs(golden, exist_ok=True)
    write_xlsx(os.path.join(golden, F_X), G, F)
    write_png(os.path.join(golden, F_PNG), F)
    with open(fig_path, "w") as fh:
        json.dump(F, fh, indent=1, sort_keys=True)
        fh.write("\n")
    with open(sub_path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(submission(F))
    return F


if __name__ == "__main__":
    a = sys.argv[1:]
    main(*a) if a else main()
