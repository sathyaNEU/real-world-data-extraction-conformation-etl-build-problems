"""task164 pack: every file the replacement program pulled on Monday 5 April 2027, written from the records.

    Bundle(T)      assembles every record from the clean truth (deterministic)
    build_pack(out, B)   writes the files into out"""
from __future__ import annotations

import csv as _csv
import datetime as dt
import io as _io
import os

import numpy as np
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import asks as AK
import common as K
import forward as FW
import postprocess as PP
import records as R
import world as W

F_TLM = {2024: "tlm_summer_peaks_2024.csv", 2025: "tlm_summer_peaks_2025.csv", 2026: "tlm_summer_peaks_2026.csv"}
F_GIS = "gis_service_transformers.csv"
F_CHG = "gis_unit_changes_2024-2027.csv"
F_SP = "service_points.csv"
F_RATE = "rate_schedule_changes.csv"
F_DMV = "dmv_pev_notifications.csv"
F_LSI = "ev_load_study_2026_intervals.parquet"
F_LSP = "ev_load_study_2026_participants.csv"
F_WX = "district_weather_daily_max.csv"
F_OMS = "oms_transformer_outages_summer2026.csv"
F_WO = "ptr_pilot_work_orders.xlsx"
F_COST = "ptr_pilot_cost_lines.csv"
F_DIST = "district_codes.csv"
F_MAN = "export_manifest.csv"
F_IEPR = "iepr_2026_pev_forecast_by_district.xlsx"
F_FDR = "feeder_peak_summary_2026.csv"
F_DICT = "data_dictionary.xlsx"
F_DEC = "decision_26-11-019_excerpt.pdf"
F_TAR = "residential_rate_schedules_excerpt.pdf"

RUN = {2024: dt.date(2024, 10, 7), 2025: dt.date(2025, 10, 6), 2026: dt.date(2026, 10, 19)}
AS_OF = pd.Timestamp(K.AS_OF)
REC = pd.Timestamp(K.RECORD_DATE)


def write_csv(path, df):
    buf = _io.StringIO()
    df.to_csv(buf, index=False, lineterminator="\n", quoting=_csv.QUOTE_MINIMAL)
    with open(path, "w", encoding="utf-8", newline="") as f:
        f.write(buf.getvalue())
    return len(df)


def feeder_label(code, f):
    return "%s-%02d%02d" % (code, f // 4 + 1, 1 + f % 4)


# ------------------------------------------------------------------------------------------------- the bundle
class Bundle:
    """Every record the pack is written from, assembled once from the truth."""

    def __init__(self, T):
        self.T = T
        tx = T.tx
        self.loc = dict(zip(tx["ti"], tx["loc"]))
        self.A = R.addresses(T)
        self.SP = R.ids(T)
        self.S = R.service_points(T, self.A, self.SP)
        self.O = AK.outages(T)
        self.O = AK.nudge_outages(T, self.O)
        self.C = AK.unit_changes(T)
        self.wos = AK.work_orders(T, self.C)
        self.C = AK.other_order_numbers(T, self.C)
        self.L = AK.cost_lines(T, self.wos)
        self.n_repl = T.pilot.groupby("district").size().to_dict()
        self.L = AK.nudge_costs(self.L, self.wos, self.n_repl)
        self.rate = R.rate_changes(T, self.SP)
        self.deep = self.deep_transformers()
        self.D = self.dmv()
        self.LP, self.LI = R.load_study(T)
        self.WX = R.weather()
        self.gis_kva = self.gis_history()

    # transformers far from the line under every member of the family (used to keep the matcher's misses and the
    # April registrations neutral)
    def deep_transformers(self, margin=9.0):
        T = self.T
        tx = T.tx
        ms = []
        for tab in T.fam.values():
            f = np.array([FW.forecast(tab, s, a, b, c, d, p) for s, a, b, c, d, p in
                          zip(tx["sg"], tx["k2026"], tx["e2026"], T.k1, T.e1, tx["pk2026"])])
            ms.append(f - K.THRESH * tx["kva"].values)
        ms = np.array(ms)
        ok = np.all(ms > margin, axis=0) | np.all(ms < -margin, axis=0)
        return set(tx["ti"].values[ok])

    def dmv(self):
        T = self.T
        r = K.rng("dmv-match")
        v = T.veh
        deep_v = v[v.ti.isin(self.deep)]
        miss = set(r.choice(deep_v.index.values, size=int(0.031 * len(v)), replace=False).tolist())
        self.unmatched_v = miss
        D = R.dmv(T, self.A, self.SP, miss)
        # registrations after the record date that were in the feed by the pull
        ra = K.rng("dmv-april")
        have = set(v["pi"].astype(int))
        pti = dict(zip(T.pr["pi"].astype(int), T.pr["ti"].astype(int)))
        free = [p for p in T.pr["pi"].values if int(p) not in have and pti[int(p)] in self.deep]
        newp = ra.choice(np.array(free), size=31, replace=False)
        rows = []
        for p in newp:
            a = self.A.loc[p]
            bev = ra.random() < .8
            mk = R.MAKES_BEV if bev else R.MAKES_PHEV
            make, model, vds, _ = mk[int(ra.integers(len(mk)))]
            my = int(ra.choice([2026, 2027, 2027]))
            vin = "%s%d%s%s%s" % (vds, int(ra.integers(0, 10)), R.YEARCODE[my], "ABCFNP"[int(ra.integers(6))], "******")
            d = dt.date(2027, 4, int(ra.integers(1, 5)))
            rows.append((vin, make, model, my, "BEV" if bev else "PHEV", R.dmv_line(a, ra), a.city.upper(), a.zip,
                         d.strftime("%Y-%m-%d"), "CURRENT", "", str(int(self.SP[p])), int(p)))
        Apr = pd.DataFrame(rows, columns=["vin_masked", "make", "model", "model_year", "fuel_type", "reg_address",
                                          "reg_city", "reg_zip", "reg_start_date", "reg_status", "status_date",
                                          "matched_sp_id", "_pi"])
        Apr["match_method"] = "EXACT"
        Apr["_vi"] = -1
        Out = R.dmv_outside(T)
        Out["match_method"] = ""
        Out["_vi"] = -1
        base = D.drop(columns=["notification_id"])
        allv = pd.concat([base, Apr, Out], ignore_index=True)
        allv["_k"] = K.rng("dmv-order").random(len(allv))
        allv = allv.sort_values(["reg_start_date", "_k"]).drop(columns="_k").reset_index(drop=True)
        rn = K.rng("dmv-ids")
        allv.insert(0, "notification_id", ["PEV%08d" % (31000000 + 3 * i + int(rn.integers(0, 3))) for i in range(len(allv))])
        self.april = set(int(p) for p in newp)
        return allv

    def gis_history(self):
        """The GIS rating of every location on each TLM run date and at the pull."""
        T = self.T
        tx = T.tx
        k0 = T.kva_orig.copy()
        k0.index = tx["ti"].values
        k0 = k0.astype(float)
        k0[tx["ti"][tx["gis_err"]].values] = tx.loc[tx["gis_err"], "kva_gis0"].values
        corr = T.other[T.other.kind == "CORR"]
        k0[corr["ti"].values] = corr["kva_from"].values
        C = self.C.sort_values(["posted", "date"])
        out = {}
        for y, d in list(RUN.items()) + [("pull", AS_OF.date())]:
            k = k0.copy()
            c = C[C["posted"] <= pd.Timestamp(d)]
            for t, kv in zip(c["ti"], c["kva_to"]):
                k[t] = kv
            out[y] = k
        return out


# ------------------------------------------------------------------------------------------------- spreadsheets
HEAD_FILL = PatternFill("solid", fgColor="DDE5EE")
BOLD = Font(name="Calibri", bold=True, size=11)


def _sheet(ws, head, rows, widths=None, fmt=None):
    ws.append(head)
    for c in ws[1]:
        c.font = BOLD
        c.fill = HEAD_FILL
    for r in rows:
        ws.append(list(r))
    for i, h in enumerate(head):
        ws.column_dimensions[get_column_letter(i + 1)].width = (widths[i] if widths else max(10, len(h) + 2))
    ws.freeze_panes = "A2"
    if fmt:
        for col, nf in fmt.items():
            for cell in ws[col][1:]:
                cell.number_format = nf


# ------------------------------------------------------------------------------------------------- writers
def tlm_frame(B, y):
    T = B.T
    tx = T.tx
    r = K.rng("tlm-ami", y)
    u = r.random(len(tx))
    ami = np.where(u < .93, 100.0, np.where(u < .99, np.round(r.uniform(99.0, 99.9, len(tx)), 1),
                                            np.round(r.uniform(96.0, 99.0, len(tx)), 1)))
    pkd = pd.Timestamp(K.summer_start(y)) + pd.to_timedelta(tx["pkday%d" % y].astype(int), unit="D")
    q = tx["pkq%d" % y].astype(int).values
    kva = B.gis_kva[y].reindex(tx["ti"]).values
    fd = [feeder_label(c, f) for c, f in zip(tx["code"], tx["feeder"])]
    if y == 2024:
        beg = [("%02d:%02d" % divmod(16 * 60 + 15 * k, 60)) for k in q]
        df = pd.DataFrame({"LOC_NO": ["%08d" % x for x in tx["loc"]], "DIST": tx["district"].str.upper().values,
                           "FDR": fd, "KVA_RTG": kva, "NUM_SP": tx["n_prem"].values,
                           "PEAK_KVA": np.round(tx["pk2024"].values, 2), "PEAK_DT": pkd.dt.strftime("%m/%d/%Y").values,
                           "PEAK_INTV_BEG": beg, "AMI_PCT": ami, "RUN_DT": RUN[y].strftime("%m/%d/%Y")})
        return df.sort_values("LOC_NO").reset_index(drop=True)
    end = [("%02d:%02d" % divmod(16 * 60 + 15 * (k + 1), 60)) for k in q]
    df = pd.DataFrame({"xfmr_loc_id": tx["loc"].values, "district_code": tx["code"].values, "feeder_id": fd,
                       "kva_rating": kva, "service_points": tx["n_prem"].values,
                       "peak_kva": np.round(tx["pk%d" % y].values, 2), "peak_date": pkd.dt.strftime("%Y-%m-%d").values,
                       "peak_interval_ending": end, "ami_read_pct": ami, "run_date": RUN[y].isoformat()})
    return df.sort_values("xfmr_loc_id").reset_index(drop=True)


def gis_frame(B):
    T = B.T
    tx = T.tx
    r = K.rng("gis-register")
    last_new = B.C[B.C["installed_serial"] != B.C["removed_serial"]].groupby("ti")["date"].max()
    inst = []
    for t in tx["ti"]:
        if t in last_new.index:
            inst.append(last_new[t].strftime("%Y-%m-%d"))
        else:
            y = T.unit_year0[int(t)]
            if y < 1986 and r.random() < .5:
                inst.append("1900-01-01")
            elif r.random() < .06:
                inst.append("")
            else:
                inst.append(dt.date(int(y), int(r.integers(1, 13)), int(r.integers(1, 29))).isoformat())
    df = pd.DataFrame({"xfmr_loc_id": ["T%d" % x for x in tx["loc"]], "district": tx["district"].values,
                       "feeder_id": [feeder_label(c, f) for c, f in zip(tx["code"], tx["feeder"])],
                       "mount": np.where(tx["mount"] == "PAD", "PADMOUNT", "OVERHEAD"),
                       "phase": r.choice(["A", "B", "C"], size=len(tx)),
                       "kva_rating": B.gis_kva["pull"].reindex(tx["ti"]).values,
                       "unit_serial": [T.serial_now[int(t)] for t in tx["ti"]],
                       "unit_installed": inst, "status": "IN SERVICE"})
    return df.sort_values("xfmr_loc_id").reset_index(drop=True)


def changes_frame(B):
    C = B.C.sort_values(["posted", "date", "ti"]).reset_index(drop=True)
    r = K.rng("change-ids")
    ids = 3304100 + np.cumsum(r.integers(1, 4, len(C)))
    return pd.DataFrame({"change_id": ["UC%07d" % i for i in ids], "xfmr_loc_id": ["T%d" % B.loc[t] for t in C["ti"]],
                         "change_date": C["date"].dt.strftime("%Y-%m-%d"), "removed_kva": C["kva_from"],
                         "installed_kva": C["kva_to"], "removed_serial": C["removed_serial"],
                         "installed_serial": C["installed_serial"], "reason_code": C["reason"],
                         "wo_number": C["wo_number"], "posted_date": C["posted"].dt.strftime("%Y-%m-%d")})


def oms_frame(B):
    O = B.O.sort_values(["event_id", "line"]).reset_index(drop=True)
    return pd.DataFrame({"event_id": O["event_id"], "line": O["line"], "xfmr_loc_id": O["xfmr_loc_id"],
                         "device_id": O["device_id"], "outage_start_utc": O["outage_start_utc"].dt.strftime("%Y-%m-%d %H:%M"),
                         "restored_utc": O["restored_utc"].dt.strftime("%Y-%m-%d %H:%M"), "cause_code": O["cause_code"],
                         "customers_out": O["customers_out"], "dispatch_district": O["dispatch_district"]})


def cost_frame(B):
    L = B.L
    return pd.DataFrame({"internal_order": "7100-PTR25", "wo_number": L["wo_number"], "cost_element": L["cost_element"],
                         "posting_date": L["posting_date"].dt.strftime("%Y-%m-%d"), "fiscal_period": L["fiscal_period"],
                         "amount_usd": L["amount"].map(lambda a: "%.2f" % a), "document": L["document"]})


def write_wo(out, B):
    path = os.path.join(out, F_WO)
    w = B.wos.sort_values("wo_number")
    wb = Workbook()
    ws = wb.active
    ws.title = "PTR pilot orders"
    rows = [(int(x.wo_number), x.district, "T%d" % B.loc[x.ti], x.plan_from, x.plan_to, x.status,
             pd.Timestamp(x.status_date).to_pydatetime(), x.crew, x.remarks) for x in w.itertuples()]
    _sheet(ws, ["Work Order", "District", "Location", "Planned From kVA", "Planned To kVA", "Status", "Status Date",
                "Crew", "Remarks"], rows, [12, 12, 12, 16, 15, 8, 12, 8, 60], fmt={"G": "yyyy-mm-dd"})
    wb.save(path)
    PP.normalize_ooxml(path, K.PEOPLE["construction"], K.PEOPLE["construction"], "2026-05-04T15:12:00Z",
                       "2027-04-05T16:47:00Z", excel=True, company=K.ORG)
    return len(rows)


def write_iepr(out, B):
    """The 2026 IEPR demand forecast filing's light-duty plug-in vehicle stock by district (mid case)."""
    path = os.path.join(out, F_IEPR)
    T = B.T
    v = T.veh
    at = pd.Timestamp("2025-12-31")
    base = v[(v.reg <= at) & (v.dep.isna() | (v.dep > at))].groupby("di").size()
    r = K.rng("iepr")
    wb = Workbook()
    ws = wb.active
    ws.title = "LDV PEV stock"
    rows = []
    for di, d in enumerate(K.DISTRICTS):
        b = float(base.get(di, 0)) * float(r.uniform(0.96, 1.03))
        g = 0.19 + 0.05 * d[6] / 0.02 + float(r.normal(0, 0.012))
        vals = [int(round(b * (1 + g) ** k / 10.0) * 10) for k in range(1, 6)]
        rows.append([d[1], d[0]] + vals)
    _sheet(ws, ["District Code", "District", "YE 2026", "YE 2027", "YE 2028", "YE 2029", "YE 2030"], rows,
           [13, 13, 10, 10, 10, 10, 10])
    ws2 = wb.create_sheet("Notes")
    for line in [["2026 IEPR demand forecast, transportation electrification. Mid case, light-duty plug-in electric vehicles in"
                  " operation at year end, by operating district (rounded to 10)."],
                 ["Base year from DMV registrations at 31 December 2025; adoption from the 2026 IEPR statewide mid case,"
                  " allocated by district household income and housing type."],
                 ["Forecasting & Analytics, filed 22 August 2026."]]:
        ws2.append(line)
    ws2.column_dimensions["A"].width = 120
    wb.save(path)
    PP.normalize_ooxml(path, "Forecasting & Analytics", "Forecasting & Analytics", "2026-07-30T17:05:00Z",
                       "2026-08-21T22:40:00Z", excel=True, company=K.ORG)
    return len(rows)


def feeder_frame(B):
    """Summer 2026 feeder-head peaks from SCADA (residential and non-residential load together)."""
    T = B.T
    tx = T.tx
    r = K.rng("feeders")
    rows = []
    for (di, f), g in tx.groupby(["di", "feeder"]):
        code = K.DISTRICTS[di][1]
        res = float((g["plateau"] * 0.82).sum()) / 1000.0
        mw = res * float(r.uniform(1.25, 1.9)) + float(r.uniform(0.4, 2.2))
        rated = round(mw / float(r.uniform(0.58, 0.93)), 1)
        d = K.HEAT[2026][0] + dt.timedelta(days=int(r.integers(0, K.HEAT[2026][1])))
        rows.append((feeder_label(code, f), K.NAMES[di], round(mw, 2), d.isoformat(), int(r.choice([17, 18, 18, 19, 19, 20])),
                     rated, round(100 * mw / rated, 1)))
    return pd.DataFrame(rows, columns=["feeder_id", "district", "peak_mw", "peak_date", "peak_hour_ending", "rating_mw",
                                       "loading_pct"])


def district_frame():
    hq = {"Foothill": "Cedar Gap", "Railyard": "Calloway (Depot St)", "Canal": "Calloway (Canal Yard)",
          "Mission": "Loma Pardo", "West Bluff": "Tule Bend", "Crestview": "Cottrell", "Lakeside": "Lake Vardon",
          "Sunrise": "Juniper Springs", "Sandhill": "Garnet Wells", "Kingsley": "Kingsley", "Highland": "Copper Flat",
          "Brookside": "Calloway (Brookside)"}
    return pd.DataFrame([(d[1], d[0], hq[d[0]], R.STATIONS[d[0]][0]) for d in K.DISTRICTS],
                        columns=["district_code", "district", "service_center", "weather_station_id"])


def write_parquet(path, I):
    ts = pa.array(I["interval_ending"].values.astype("datetime64[s]"), type=pa.timestamp("s"))
    t = pa.Table.from_arrays([pa.array(I["participant_id"].values, type=pa.string()), ts,
                              pa.array(I["kwh"].values.astype("float32"), type=pa.float32())],
                             names=["participant_id", "interval_ending", "kwh"])
    pq.write_table(t, path, compression="zstd", row_group_size=500000, use_dictionary=["participant_id"])
    return len(I)


def build_pack(out, B):
    import docs as DOC
    os.makedirs(out, exist_ok=True)
    p = lambda f: os.path.join(out, f)
    n = {}
    for y in K.SUMMERS:
        n[F_TLM[y]] = write_csv(p(F_TLM[y]), tlm_frame(B, y))
    n[F_GIS] = write_csv(p(F_GIS), gis_frame(B))
    n[F_CHG] = write_csv(p(F_CHG), changes_frame(B))
    n[F_SP] = write_csv(p(F_SP), B.S)
    n[F_RATE] = write_csv(p(F_RATE), B.rate)
    n[F_DMV] = write_csv(p(F_DMV), B.D.drop(columns=["_pi", "_vi"]))
    n[F_LSI] = write_parquet(p(F_LSI), B.LI)
    n[F_LSP] = write_csv(p(F_LSP), B.LP)
    n[F_WX] = write_csv(p(F_WX), B.WX)
    n[F_OMS] = write_csv(p(F_OMS), oms_frame(B))
    n[F_WO] = write_wo(out, B)
    n[F_COST] = write_csv(p(F_COST), cost_frame(B))
    n[F_DIST] = write_csv(p(F_DIST), district_frame())
    n[F_IEPR] = write_iepr(out, B)
    n[F_FDR] = write_csv(p(F_FDR), feeder_frame(B))
    DOC.write_all(out, B)
    n[F_MAN] = write_csv(p(F_MAN), manifest(n))
    return n


def manifest(n):
    rows = [
        (F_TLM[2024], "Transformer load management summer peaks, 2024 season run (TLM 4.2 layout)", "TLM 4.2", "Distribution Planning", "2024-10-07", n[F_TLM[2024]]),
        (F_TLM[2025], "Transformer load management summer peaks, 2025 season run", "TLM 5", "Distribution Planning", "2025-10-06", n[F_TLM[2025]]),
        (F_TLM[2026], "Transformer load management summer peaks, 2026 season run", "TLM 5", "Distribution Planning", "2026-10-19", n[F_TLM[2026]]),
        (F_GIS, "Single-phase residential service transformer locations in service, as at the pull", "GIS", "GIS Records", "2027-04-05", n[F_GIS]),
        (F_CHG, "Unit changes posted to GIS, 1 January 2024 to the pull", "GIS", "GIS Records", "2027-04-05", n[F_CHG]),
        (F_SP, "Residential service points with their transformer, as at the pull", "CIS", "Customer Systems", "2027-04-05", n[F_SP]),
        (F_RATE, "Rate schedule changes to R-EV2, 2018 to the pull", "CIS", "Customer Systems", "2027-04-05", n[F_RATE]),
        (F_DMV, "DMV plug-in vehicle notifications with the address match, 2018 to the pull", "DMV PEV feed", "Customer Systems", "2027-04-05", n[F_DMV]),
        (F_LSI, "EV charging load study, EVSE submeter intervals, summer 2026", "Load Research", "Load Research", "2026-11-16", n[F_LSI]),
        (F_LSP, "EV charging load study participants", "Load Research", "Load Research", "2026-11-16", n[F_LSP]),
        (F_WX, "District weather stations, daily maximum temperature, summers 2024 to 2026", "Met network", "Load Research", "2027-04-05", n[F_WX]),
        (F_OMS, "Outage management system, transformer outage lines, 16 May to 15 October 2026", "OMS", "Reliability", "2026-10-16", n[F_OMS]),
        (F_WO, "2025-26 proactive replacement pilot work orders", "Work management", "Construction Services", "2027-04-05", n[F_WO]),
        (F_COST, "Internal order 7100-PTR25 line items to the pull", "General ledger", "Finance", "2027-04-05", n[F_COST]),
        (F_DIST, "District codes, service centers and weather stations", "Reference", "Distribution Planning", "2027-04-05", n[F_DIST]),
        (F_IEPR, "2026 IEPR plug-in vehicle forecast by district", "Forecasting", "Forecasting & Analytics", "2026-08-22", n[F_IEPR]),
        (F_FDR, "Feeder-head peaks, summer 2026", "SCADA historian", "Distribution Planning", "2026-10-19", n[F_FDR]),
        (F_DICT, "Field definitions for this pull", "", "Regulatory Compliance", "2027-04-05", ""),
        (F_DEC, "Decision 26-11-019, excerpt", "", "Regulatory Compliance", "2026-11-12", ""),
        (F_TAR, "Residential rate schedules, excerpt", "", "Regulatory Compliance", "2026-03-01", ""),
    ]
    return pd.DataFrame(rows, columns=["file", "contents", "source_system", "owner", "extract_date", "rows"])
