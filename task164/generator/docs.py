"""task164 documents: the Commission's decision excerpt, the residential rate schedules excerpt and the data
dictionary for the pull. Their text is kept here as data so checks.py and verify.py can quote it."""
from __future__ import annotations

import datetime as dt
import os

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                        # noqa: E402
from reportlab.lib.pagesizes import letter                                              # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                         # noqa: E402
from reportlab.lib.units import mm                                                      # noqa: E402
from reportlab.pdfbase import pdfmetrics                                                # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                            # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402
from openpyxl import Workbook                                                           # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill                                # noqa: E402

import common as K          # noqa: E402
import postprocess as PP    # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/"
for _n, _f in (("LibSerif", "liberation/LiberationSerif-Regular.ttf"), ("LibSerif-Bold", "liberation/LiberationSerif-Bold.ttf"),
               ("LibSans", "liberation/LiberationSans-Regular.ttf"), ("LibSans-Bold", "liberation/LiberationSans-Bold.ttf")):
    pdfmetrics.registerFont(TTFont(_n, FONT_DIR + _f))
INK = colors.HexColor("#141414")
GREY = colors.HexColor("#555555")
RULE = colors.HexColor("#9a9a9a")
HEAD = colors.HexColor("#e9edf2")

F_DEC = "decision_26-11-019_excerpt.pdf"
F_TAR = "residential_rate_schedules_excerpt.pdf"
F_DICT = "data_dictionary.xlsx"


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def pdf_stamp(d, hh=9, mi=0):
    off = 7 if dt.date(d.year, 3, 8) <= d < dt.date(d.year, 11, 1) else 8
    return ("D:%04d%02d%02d%02d%02d00-%02d'00'" % (d.year, d.month, d.day, hh, mi, off)).encode()


def styles(font, bold, size):
    return dict(
        body=ParagraphStyle("b" + font, fontName=font, fontSize=size, leading=size * 1.38, textColor=INK, spaceAfter=6),
        center=ParagraphStyle("c" + font, fontName=bold, fontSize=size + 0.6, leading=size * 1.4, textColor=INK,
                              alignment=1, spaceAfter=4),
        h2=ParagraphStyle("h2" + font, fontName=bold, fontSize=size + 0.4, leading=size * 1.5, textColor=INK,
                          spaceBefore=6, spaceAfter=4),
        small=ParagraphStyle("sm" + font, fontName=font, fontSize=size - 1.6, leading=(size - 1.6) * 1.35, textColor=GREY,
                             spaceAfter=4),
        cell=ParagraphStyle("ce" + font, fontName=font, fontSize=size - 1.4, leading=(size - 1.4) * 1.25, textColor=INK),
        cellb=ParagraphStyle("cb" + font, fontName=bold, fontSize=size - 1.4, leading=(size - 1.4) * 1.25, textColor=INK))


def _doc(path, title, author, margin=24):
    return SimpleDocTemplate(path, pagesize=letter, leftMargin=margin * mm, rightMargin=margin * mm, topMargin=20 * mm,
                             bottomMargin=19 * mm, title=title, author=author, subject="", creator="", producer="")


def _footer(left, font):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 8.0)
        canvas.setFillColor(GREY)
        canvas.drawString(24 * mm, 11 * mm, left)
        canvas.drawRightString(letter[0] - 24 * mm, 11 * mm, "- %d -" % doc.page)
        canvas.restoreState()
    return draw


# ------------------------------------------------------------------------------------------------- the decision
DEC_DATE = K.DECISION_DATE
DEC_HEAD = [
    "Decision 26-11-019  November 12, 2026",
    "Application 26-03-008",
    "Application of Calloway Valley Power Company for Approval of Its 2027 Proactive Service Transformer Replacement Program",
    "DECISION APPROVING 2027 PROACTIVE SERVICE TRANSFORMER REPLACEMENTS",
]
DEC_FINDINGS = [
    "4. Calloway Valley Power Company (Calloway Valley Power) serves its residential customers through about 51,800 "
    "single-phase residential service transformers in twelve operating districts: Foothill, Railyard, Canal, Mission, "
    "West Bluff, Crestview, Lakeside, Sunrise, Sandhill, Kingsley, Highland and Brookside.",
    "5. Calloway Valley Power's transformer load management system sums, for each service transformer, the interval meter "
    "data of the service points the transformer serves, and records for each summer (June 1 through September 30) the "
    "transformer's highest quarter-hour demand between 4 p.m. and 9 p.m., its recorded summer peak.",
    "6. Under Decision 25-06-011 Calloway Valley Power replaced overloaded service transformers in the Railyard, Canal "
    "and Brookside districts between October 2025 and April 2026.",
    "7. A service transformer whose summer peak exceeds 120 percent of its nameplate rating loses insulation life "
    "rapidly in a heat event and is at materially higher risk of failure.",
    "8. Calloway Valley Power proposes to replace 400 service transformers in 2027 ahead of failure and to direct the "
    "replacements to the districts where transformers are most at risk.",
]
DEC_ORDERS = [
    "1. Calloway Valley Power Company is authorized to replace 400 single-phase residential service transformers in "
    "calendar year 2027 under its Proactive Service Transformer Replacement Program.",
    "2. The 400 replacements shall be allocated among Calloway Valley Power's twelve operating districts in proportion to "
    "each district's forecast count of transformers at risk. Each district shall first receive the whole-number part of "
    "its proportional share, and the replacements remaining shall then be assigned one each to the districts with the "
    "largest fractional parts.",
    "3. A transformer is at risk if its forecast 2027 summer peak exceeds 120 percent of its nameplate rating. Every "
    "single-phase residential service transformer in service on the record date shall be forecast, and each shall be "
    "counted in the operating district it belongs to on that date.",
    "4. A transformer's forecast 2027 summer peak is its recorded 2026 summer peak adjusted for the difference between "
    "the plug-in electric vehicles registered at its service addresses on the day of its recorded 2026 peak and those "
    "registered there on the record date, each at the rate schedule in effect for its service address on the date "
    "concerned. Each such vehicle shall be taken at the expected effect of one plug-in electric vehicle on a service "
    "transformer's recorded summer peak, one figure applying to every vehicle on Schedule R-EV2 and one to every vehicle "
    "on any other residential schedule. No other change in load shall be assumed, and no vehicle registered after the "
    "record date shall be counted.",
    "5. The record date is March 31, 2027. A transformer's nameplate rating is that of the unit installed at its "
    "location on the record date.",
    "6. Calloway Valley Power shall file a Tier 2 advice letter by April 30, 2027, setting out each district's forecast "
    "count of transformers at risk and its allocation.",
    "7. Application 26-03-008 remains open.",
]


def write_decision(out):
    st = styles("LibSerif", "LibSerif-Bold", 11.0)
    path = os.path.join(out, F_DEC)
    doc = _doc(path, "Decision 26-11-019 (excerpt)", "Regulatory Compliance")
    flow = [Paragraph(esc(DEC_HEAD[0]), st["small"]), Paragraph(esc(DEC_HEAD[1]), st["small"]), Spacer(1, 6),
            Paragraph(esc(DEC_HEAD[2]), st["center"]), Spacer(1, 4), Paragraph(esc(DEC_HEAD[3]), st["center"]),
            Spacer(1, 4),
            Paragraph(esc("Excerpt: Findings of Fact 4 to 8 and the Ordering Paragraphs, as served November 12, 2026."),
                      st["small"]),
            Paragraph("Findings of Fact", st["h2"])]
    flow += [Paragraph(esc(p), st["body"]) for p in DEC_FINDINGS]
    flow.append(Paragraph("ORDER", st["h2"]))
    flow.append(Paragraph("IT IS ORDERED that:", st["body"]))
    flow += [Paragraph(esc(p), st["body"]) for p in DEC_ORDERS]
    flow.append(Spacer(1, 8))
    flow.append(Paragraph(esc("This order is effective today. Dated November 12, 2026."), st["body"]))
    foot = "A.26-03-008  Decision 26-11-019"
    doc.build(flow, onFirstPage=_footer(foot, "LibSerif"), onLaterPages=_footer(foot, "LibSerif"))
    PP.normalize_pdf(path, pdf_stamp(DEC_DATE, 16, 2), pdf_stamp(DEC_DATE, 16, 2))


# ------------------------------------------------------------------------------------------------- the tariffs
TAR_DATE = dt.date(2026, 3, 1)
TARIFF = [
    ("Schedule R-TOU4. Residential Time-of-Use (Peak Pricing 4 p.m. to 9 p.m. Every Day)", [
        "Applicability. The default schedule for residential service. Customers enrolled in the CARE program receive the "
        "CARE discount on this schedule.",
        "Time periods, every day of the year: Peak 4:00 p.m. to 9:00 p.m.; Off-Peak all other hours.",
    ], [("Summer (June 1 to September 30)", "0.548", "0.437"), ("Winter (October 1 to May 31)", "0.462", "0.418")],
        ("Season", "Peak $/kWh", "Off-Peak $/kWh")),
    ("Schedule R-EV2. Residential Time-of-Use, Plug-In Electric Vehicle", [
        "Applicability. Whole-premises residential service for a customer with a plug-in electric vehicle registered at "
        "the service address. Customers enrolled in the CARE program receive the CARE discount on this schedule.",
        "Time periods, every day of the year: Peak 4:00 p.m. to 9:00 p.m.; Partial-Peak 3:00 p.m. to 4:00 p.m. and "
        "9:00 p.m. to 12:00 midnight; Off-Peak 12:00 midnight to 3:00 p.m.",
    ], [("Summer (June 1 to September 30)", "0.612", "0.507", "0.309"), ("Winter (October 1 to May 31)", "0.481", "0.465", "0.301")],
        ("Season", "Peak $/kWh", "Partial-Peak $/kWh", "Off-Peak $/kWh")),
    ("Schedule R-1. Residential Service (closed to new customers)", [
        "Applicability. Tiered residential service, open only to customers served on it on March 1, 2021.",
    ], [("Baseline quantity", "0.356"), ("101 to 400 percent of baseline", "0.447"), ("Above 400 percent of baseline", "0.536")],
        ("Usage", "$/kWh")),
]
TAR_SPECIAL = [
    "Changes of schedule. A customer may change residential schedules once in any twelve months. A change takes effect "
    "on the customer's next regular meter reading date after the request is received.",
]


def write_tariffs(out):
    st = styles("LibSans", "LibSans-Bold", 9.6)
    path = os.path.join(out, F_TAR)
    doc = _doc(path, "Residential rate schedules (excerpt)", "Regulatory Compliance", margin=21)
    flow = [Paragraph(esc("Calloway Valley Power Company. Electric rate schedules, residential (excerpt). Rates "
                          "effective March 1, 2026."), st["h2"]), Spacer(1, 2)]
    for head, paras, rows, cols in TARIFF:
        flow.append(Paragraph(esc(head), st["h2"]))
        flow += [Paragraph(esc(p), st["body"]) for p in paras]
        data = [[Paragraph(esc(c), st["cellb"]) for c in cols]] + [[Paragraph(esc(x), st["cell"]) for x in r] for r in rows]
        t = Table(data, colWidths=[62 * mm] + [30 * mm] * (len(cols) - 1))
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), HEAD), ("LINEBELOW", (0, 0), (-1, 0), 0.5, RULE),
                               ("LINEBELOW", (0, -1), (-1, -1), 0.5, RULE), ("TOPPADDING", (0, 0), (-1, -1), 2),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
        flow += [t, Spacer(1, 5)]
    flow.append(Paragraph("Special conditions", st["h2"]))
    flow += [Paragraph(esc(p), st["body"]) for p in TAR_SPECIAL]
    flow.append(Paragraph(esc("Energy charges only; customer charges, the CARE discount and non-bypassable charges are "
                              "set out in the full schedules."), st["small"]))
    foot = "Calloway Valley Power Company  Residential rate schedules (excerpt)"
    doc.build(flow, onFirstPage=_footer(foot, "LibSans"), onLaterPages=_footer(foot, "LibSans"))
    PP.normalize_pdf(path, pdf_stamp(TAR_DATE, 10, 14), pdf_stamp(TAR_DATE, 10, 14))


# ------------------------------------------------------------------------------------------------- the dictionary
DICT = {
    "tlm_summer_peaks_2025/2026": [
        ("xfmr_loc_id", "GIS location number of the service transformer (the number after the T)."),
        ("district_code", "Operating district code (district_codes.csv)."),
        ("feeder_id", "Distribution feeder."),
        ("kva_rating", "Nameplate kVA of the unit in GIS on the run date."),
        ("service_points", "Service points connected to the transformer on the run date."),
        ("peak_kva", "Highest quarter-hour demand of the transformer's summed service points from 1 June to 30 September, "
                     "among the intervals ending 16:15 to 21:00 local time. Interval kWh are converted to kW and then to "
                     "kVA at a power factor of 0.95."),
        ("peak_date", "Day of peak_kva."),
        ("peak_interval_ending", "End of the quarter hour of peak_kva, local time."),
        ("ami_read_pct", "Share of the summer's interval reads received for the transformer's service points."),
        ("run_date", "Date of the season run."),
    ],
    "tlm_summer_peaks_2024": [
        ("LOC_NO", "GIS location number, eight digits with leading zeros."),
        ("DIST", "Operating district."),
        ("FDR", "Distribution feeder."),
        ("KVA_RTG", "Nameplate kVA of the unit in GIS on the run date."),
        ("NUM_SP", "Service points connected on the run date."),
        ("PEAK_KVA", "Highest quarter-hour demand of the summed service points from 1 June to 30 September, among the "
                     "intervals beginning 16:00 to 20:45 local time, in kVA at a power factor of 0.95."),
        ("PEAK_DT", "Day of PEAK_KVA (MM/DD/YYYY)."),
        ("PEAK_INTV_BEG", "Start of the quarter hour of PEAK_KVA, local time."),
        ("AMI_PCT", "Share of the summer's interval reads received."),
        ("RUN_DT", "Date of the season run (MM/DD/YYYY)."),
    ],
    "gis_service_transformers": [
        ("xfmr_loc_id", "Transformer location."),
        ("district", "Operating district."),
        ("feeder_id", "Distribution feeder."),
        ("mount", "OVERHEAD (pole) or PADMOUNT."),
        ("phase", "Primary phase the transformer is connected to."),
        ("kva_rating", "Nameplate kVA of the installed unit."),
        ("unit_serial", "Manufacturer serial of the installed unit."),
        ("unit_installed", "Date the installed unit was set. 1900-01-01 where the paper records were not converted."),
        ("status", "Location status."),
    ],
    "gis_unit_changes_2024-2027": [
        ("change_id", "GIS change."),
        ("xfmr_loc_id", "Transformer location."),
        ("change_date", "Date of the change in the field."),
        ("removed_kva", "GIS kVA rating at the location before the change."),
        ("installed_kva", "GIS kVA rating at the location after the change."),
        ("removed_serial", "Unit serial at the location before the change."),
        ("installed_serial", "Unit serial at the location after the change."),
        ("reason_code", "PTR proactive transformer replacement program; FAIL failed unit; MAINT condition (tank, bushing, "
                        "leak); LOAD customer load; CORR records."),
        ("wo_number", "Work order the change was posted from, where recorded."),
        ("posted_date", "Date the change was posted to GIS."),
    ],
    "service_points": [
        ("sp_id", "Service point."),
        ("xfmr_loc_id", "Transformer serving the service point."),
        ("service_address", "Service address."),
        ("city", "City."),
        ("zip", "ZIP code."),
        ("district", "Operating district."),
        ("rate_schedule", "Rate schedule on the pull date."),
        ("care", "Y if enrolled in CARE on the pull date."),
        ("account_start", "Start of the current account; blank where not converted from the old billing system."),
    ],
    "rate_schedule_changes": [
        ("change_id", "Schedule change."),
        ("sp_id", "Service point."),
        ("from_schedule", "Schedule before the change."),
        ("to_schedule", "Schedule after the change."),
        ("request_date", "Date the customer's request was received."),
        ("effective_date", "First day billed on the new schedule."),
        ("channel", "How the request came in: WEB, CONTACT CTR, DEALER REF (dealer referral), CAMPAIGN (the winter "
                    "2026-27 R-EV2 enrollment campaign)."),
    ],
    "dmv_pev_notifications": [
        ("notification_id", "Notification from the DMV plug-in vehicle feed."),
        ("vin_masked", "First eleven characters of the VIN."),
        ("make", "Make."), ("model", "Model."), ("model_year", "Model year."),
        ("fuel_type", "BEV battery electric; PHEV plug-in hybrid."),
        ("reg_address", "Registered address as DMV holds it."),
        ("reg_city", "Registered city."), ("reg_zip", "Registered ZIP code."),
        ("reg_start_date", "Date the vehicle was registered at this address."),
        ("reg_status", "CURRENT; EXPIRED (registration lapsed); TRANSFERRED (sold); ADDR CHANGE (owner moved); OUT OF "
                       "STATE."),
        ("status_date", "Date of reg_status; blank for CURRENT."),
        ("matched_sp_id", "Service point the utility's address match placed the registration at; blank where the "
                          "address did not match."),
        ("match_method", "EXACT or NORMALIZED (matched after standardizing the address)."),
    ],
    "ev_load_study_2026_participants": [
        ("participant_id", "Study participant."),
        ("district", "Operating district."),
        ("rate_schedule", "Rate schedule throughout the study."),
        ("fuel_type", "BEV or PHEV."),
        ("evse_level", "L1 (120 V) or L2 (240 V)."),
        ("charging_kw", "Charging rate measured at the EVSE, kW."),
        ("data_start", "First day of interval data."),
        ("intervals_expected", "Quarter-hour intervals from data_start to 30 September 2026."),
        ("intervals_received", "Intervals received."),
    ],
    "ev_load_study_2026_intervals": [
        ("participant_id", "Study participant."),
        ("interval_ending", "End of the quarter hour, local time."),
        ("kwh", "Energy delivered by the EVSE in the quarter hour, kWh. Intervals with no read are absent."),
    ],
    "district_weather_daily_max": [
        ("station_id", "Weather station (district_codes.csv)."), ("station_name", "Station."),
        ("date", "Day."), ("tmax_f", "Daily maximum temperature, degrees F."), ("qc_flag", "M missing."),
    ],
    "oms_transformer_outages_summer2026": [
        ("event_id", "Outage event."),
        ("line", "Line within the event."),
        ("xfmr_loc_id", "Transformer de-energised."),
        ("device_id", "Device that operated or was opened: CO cutout, XF transformer fuse, LF lateral fuse, EF loop "
                      "elbow fuse, followed by its number."),
        ("outage_start_utc", "Start of the outage, UTC."),
        ("restored_utc", "Time the transformer was re-energised, UTC."),
        ("cause_code", "EQ-XFMR transformer failure; OL-XFMR transformer fuse operated on overload, no failure; EQ-CUTOUT "
                       "cutout or fuse hardware; EQ-CONN connector or secondary; TREE; ANML animal; VEH vehicle; PLND "
                       "planned; UNK unknown."),
        ("customers_out", "Customers out of service."),
        ("dispatch_district", "District whose crew was dispatched. Outage reporting counts each line in the operating "
                              "district of the transformer it de-energised (gis_service_transformers)."),
    ],
    "ptr_pilot_work_orders": [
        ("Work Order", "Work order."), ("District", "Operating district."), ("Location", "Transformer location."),
        ("Planned From kVA", "GIS rating when the order was raised."), ("Planned To kVA", "Size the order planned to set."),
        ("Status", "COMP completed and closed; CANC cancelled."), ("Status Date", "Date of the status."),
        ("Crew", "Crew; CC are contract crews."), ("Remarks", "Crew's closing remarks."),
    ],
    "ptr_pilot_cost_lines": [
        ("internal_order", "Internal order of the 2025-26 pilot."),
        ("wo_number", "Work order charged."),
        ("cost_element", "LABOR company crew; CONTRACT contract crew invoice; CONTRACT-ACCR month-end accrual of contract "
                         "work, reversed the next day; MATL-XFMR transformer; MATL-OTHER fittings and hardware; FLEET."),
        ("posting_date", "Ledger posting date."),
        ("fiscal_period", "Fiscal period (calendar month)."),
        ("amount_usd", "Amount, dollars."),
        ("document", "Ledger document; -R a reversal, -C the posting that replaces it."),
    ],
    "export_manifest": [
        ("file", "File in the pull."), ("contents", "What the file holds."), ("source_system", "System it was extracted from."),
        ("owner", "Owning group."), ("extract_date", "Date of the extract."), ("rows", "Data rows."),
    ],
    "district_codes": [
        ("district_code", "Code."), ("district", "Operating district."), ("service_center", "Crew service center."),
        ("weather_station_id", "District weather station."),
    ],
    "iepr_2026_pev_forecast_by_district": [
        ("YE 2026 to YE 2030", "Light-duty plug-in vehicles in operation at year end, mid case, rounded to 10."),
    ],
    "feeder_peak_summary_2026": [
        ("feeder_id", "Feeder."), ("district", "Operating district."), ("peak_mw", "Feeder-head peak, MW (SCADA)."),
        ("peak_date", "Day of the peak."), ("peak_hour_ending", "Hour ending of the peak, local time."),
        ("rating_mw", "Feeder rating, MW."), ("loading_pct", "Peak as a share of rating."),
    ],
}
DICT_NOTES = [
    "Data dictionary for the 5 April 2027 pull, 2027 proactive transformer replacement filing.",
    "Local time is Pacific prevailing time (PDT from the second Sunday in March to the first Sunday in November). Fields "
    "ending _utc are in UTC.",
    "tlm_summer_peaks: one row per transformer location in service on the run date.",
    "gis_service_transformers: one row per transformer location in service on the pull date.",
    "gis_unit_changes: one row per change posted to GIS.",
    "service_points: one row per residential service point active on the pull date.",
    "dmv_pev_notifications: one row per vehicle registration at an address.",
    "ev_load_study_2026_intervals: one row per participant and quarter hour received.",
    "oms_transformer_outages_summer2026: one row per transformer de-energised in an outage event.",
    "ptr_pilot_cost_lines: one row per ledger line on internal order 7100-PTR25.",
]


def write_dictionary(out):
    path = os.path.join(out, F_DICT)
    wb = Workbook()
    ws = wb.active
    ws.title = "Notes"
    for line in DICT_NOTES:
        ws.append([line])
    ws.column_dimensions["A"].width = 120
    fill = PatternFill("solid", fgColor="DDE5EE")
    ws2 = wb.create_sheet("Fields")
    ws2.append(["File", "Field", "Definition"])
    for c in ws2[1]:
        c.font = Font(name="Calibri", bold=True, size=11)
        c.fill = fill
    for f, rows in DICT.items():
        for field, d in rows:
            ws2.append([f, field, d])
    ws2.column_dimensions["A"].width = 34
    ws2.column_dimensions["B"].width = 22
    ws2.column_dimensions["C"].width = 110
    for row in ws2.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
    ws2.freeze_panes = "A2"
    wb.save(path)
    PP.normalize_ooxml(path, K.PEOPLE["compliance"], K.PEOPLE["compliance"], "2027-03-15T16:20:00Z",
                       "2027-04-05T21:18:00Z", excel=True, company=K.ORG)


def write_all(out, B=None):
    write_decision(out)
    write_tariffs(out)
    write_dictionary(out)
