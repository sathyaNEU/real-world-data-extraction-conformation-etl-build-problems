"""Write the evidence pack's data files from the simulated world."""
import csv
import datetime as dt
import json

import numpy as np

import world as W

FILES = {
    "register": "complaint_matters_2025_2026.csv",
    "history": "matter_status_history_2025_2026.csv",
    "social": "social_care_posts_2025_2026.jsonl",
    "meridel": "meridel_affected_payments_20260508_10.csv",
    "portfolio": "bellcroft_card_portfolio_2025_2026.csv",
    "codes": "complaint_codes_v7.csv",
    "time": "fenwright_time_bellcroft_site_2026.csv",
    "marlowby": "marlowby_timesheets_fenwright_2026.csv",
    "treadwell": "treadwell_oake_hours_2026.xlsx",
    "opslog": "desk_operations_log_2026.xlsx",
    "manifest": "folder_contents_2026-10-09.csv",
    "ivr": "ivr_complaint_line_weekly_2026.csv",
    "cfpb": "cfpb_public_complaints_bellcroft_card.csv",
}
DISTRACTORS = ["ivr", "cfpb"]
QUEUE = {"Phone": "GEN", "Digital": "GEN", "CFPB": "REG", "State": "REG", "Executive": "REG", "Social": "SOC"}
STATES = ["OH", "IN", "KY", "PA", "MI", "WV"]


def ts(t):
    return t.strftime("%Y-%m-%d %H:%M:%S")


def _writer(path, header):
    fh = open(path, "w", newline="", encoding="utf-8")
    w = csv.writer(fh, lineterminator="\n")
    w.writerow(header)
    return fh, w


def source_refs(rng, rows):
    """Each channel's own reference for the contact that opened the matter (the social care desk's post ids are
    set when its records are built)."""
    seq = {"Phone": 0, "Digital": 0, "State": 0, "Executive": 0}
    for r in rows:
        ch = r["channel"]
        if ch == "Social":
            continue
        d = r["opened"].date()
        if ch == "Phone":
            seq[ch] += 1
            r["source_ref"] = "IVR%s%05d" % (d.strftime("%y%m%d"), (seq[ch] * 37) % 100000)
        elif ch == "Digital":
            seq[ch] += 1
            r["source_ref"] = "SM-%07d" % (4400000 + seq[ch] * 3)
        elif ch == "CFPB":
            rec = d - dt.timedelta(days=int(rng.integers(1, 9)))
            r["source_ref"] = "%s-%07d" % (rec.strftime("%y%m%d"), int(rng.integers(1000000, 9999999)))
            r["cfpb_received"] = rec
        elif ch == "State":
            seq[ch] += 1
            st = STATES[int(rng.integers(0, len(STATES)))]
            r["source_ref"] = "%s-AG-%s-%05d" % (st, d.strftime("%y"), 1000 + seq[ch] * 7 % 90000)
        else:
            seq[ch] += 1
            r["source_ref"] = "EO-%d-%04d" % (d.year, seq[ch])


def write_register(rows, path):
    fh, w = _writer(path, ["matter_no", "opened_at", "channel", "source_ref", "problem_code", "product",
                           "account_basis", "account_ref", "queue", "status", "closed_at"])
    for r in rows:
        ev = r["events"]
        last = ev[-1][1]
        status = "MERGED" if last == "MERGED" else ("CLOSED" if last == "CLOSED" else "OPEN")
        closed = ts(ev[-1][0]) if last in ("CLOSED", "MERGED") else ""
        w.writerow([r["matter_no"], ts(r["opened"]), r["channel"], r["source_ref"], r["code"], r["product"],
                    r["basis"], r["account"], QUEUE[r["channel"]], status, closed])
    fh.close()
    return len(rows)


def write_history(rows, path):
    out = []
    for r in rows:
        for t, e in r["events"]:
            out.append((t, r["matter_no"], e))
    out.sort()
    fh, w = _writer(path, ["matter_no", "event_at", "event"])
    for t, m, e in out:
        w.writerow([m, ts(t), e])
    fh.close()
    return len(out)


def write_social(records, path):
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        for rec in sorted(records, key=lambda x: (x["created_at"], x["post_id"])):
            fh.write(json.dumps(rec, ensure_ascii=False, separators=(",", ":")) + "\n")
    return len(records)


def write_meridel(payments, path):
    fh, w = _writer(path, ["payment_ref", "acct_no", "received_date", "due_post_date", "posted_date",
                           "business_days_late", "amount", "status", "return_code", "original_payment_ref"])
    for p in payments:
        w.writerow([p["payment_ref"], "00" + p["account"][3:], p["received"].isoformat(), p["due"].isoformat(),
                    p["posted"].isoformat() if p["posted"] else "", "" if p["late"] is None else p["late"],
                    "%.2f" % p["amount"], p["status"], p["return_code"], p["original_ref"]])
    fh.close()
    return len(payments)


def write_portfolio(rng, path):
    fh, w = _writer(path, ["month_end", "open_accounts", "new_accounts", "closed_accounts"])
    prev = None
    n = 0
    for d, v in sorted(W.PORTFOLIO_OPEN.items()):
        if prev is None:
            prev = v
            continue
        new = int(rng.integers(12800, 15600))
        closed = new - (v - prev)
        w.writerow([d.isoformat(), v, new, closed])
        prev = v
        n += 1
    fh.close()
    return n


def write_codes(path):
    fh, w = _writer(path, ["code", "description", "family", "in_use_from"])
    for c, d, f in W.CODES:
        w.writerow([c, d, f, "2025-10-01"])
    fh.close()
    return len(W.CODES)


def write_time(time_rows, path):
    fh, w = _writer(path, ["employee_id", "work_date", "segment_start", "segment_end", "pay_code", "program", "hours"])
    for r in time_rows:
        w.writerow([r["emp"], r["date"].isoformat(), r["start"], r["end"], r["code"], r["program"],
                    ("%.2f" % r["hours"]).rstrip("0").rstrip(".")])
    fh.close()
    return len(time_rows)


def write_marlowby(rows, path):
    fh, w = _writer(path, ["worker_no", "shift_date", "time_in", "time_out", "unpaid_break_min", "hours",
                           "assignment", "week_ending", "export_id"])
    for r in rows:
        w.writerow([r["worker"], r["date"].strftime("%m/%d/%Y"), r["tin"], r["tout"], r["brk"],
                    ("%.2f" % r["hours"]), "FENWRIGHT CONTACT SVCS - BELLCROFT CMPL", r["week_ending"].strftime("%m/%d/%Y"),
                    r["export"]])
    fh.close()
    return len(rows)


def write_treadwell(rows, path):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hours"
    ws.append(["Associate ID", "Work Date", "Start", "End", "Break (min)", "Hours", "Site", "Approved By"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r["assoc"], r["date"], r["start"], r["end"], r["brk"], r["hhmm"], "Fenwright Columbus / Bellcroft",
                   r["approver"]])
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for c in row:
            c.number_format = "m/d/yyyy"
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        for c in row:
            c.number_format = "0.00"
    for col, wd in zip("ABCDEFGH", (13, 11, 8, 8, 11, 8, 30, 12)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A2"
    wb.save(path)
    return len(rows)


def write_ivr(rng, path):
    fh, w = _writer(path, ["week_starting", "calls_offered", "calls_answered", "calls_abandoned",
                           "ASA (sec)", "AHT (sec)"])
    d = dt.date(2025, 12, 29)
    n = 0
    while d <= dt.date(2026, 9, 28):
        inc = W.INCIDENT_PHONE.get(d, 0)
        offered = int(rng.normal(3150, 140) + inc * 3.9)
        ans = int(offered * float(np.clip(0.95 - inc / 9000.0 + rng.normal(0, 0.006), 0.6, 0.99)))
        w.writerow([d.isoformat(), offered, ans, offered - ans, int(28 + inc / 9 + rng.normal(0, 4)),
                    int(rng.normal(468, 14) + inc / 25)])
        d += dt.timedelta(days=7)
        n += 1
    fh.close()
    return n


CFPB_ISSUE = {"PAY-01": ("Problem when making payments", "Problem during payment process"),
              "PAY-02": ("Problem when making payments", "You never received your bill or did not know a payment was due"),
              "FEE-01": ("Fees or interest", "Problem with fees"),
              "FEE-02": ("Fees or interest", "Charged too much interest"),
              "FEE-03": ("Fees or interest", "Problem with fees"),
              "CRB-01": ("Incorrect information on your report", "Account status incorrect"),
              "CRB-02": ("Incorrect information on your report", "Account information incorrect"),
              "DSP-01": ("Problem with a purchase shown on your statement", "Credit card company isn't resolving a dispute about a purchase on your statement"),
              "FRD-01": ("Problem with a purchase shown on your statement", "Card was charged for something you did not purchase with the card"),
              "RWD-01": ("Advertising and marketing, including promotional offers", "Didn't receive advertised or promotional terms"),
              "SVC-01": ("Getting a credit card", "Problem with the process of getting a credit card"),
              "LIM-01": ("Credit limit changed", "Credit limit changed"),
              "CLS-01": ("Closing your account", "Company closed your account"),
              "COL-01": ("Struggling to pay your bill", "Problem with a hardship program"),
              "AGT-01": ("Other features, terms, or problems", "Other problem"),
              "APP-01": ("Getting a credit card", "Card opened without my consent or knowledge"),
              "OTH-99": ("Other features, terms, or problems", "Other problem")}


def write_cfpb(rng, rows, path):
    """The CFPB's published complaints about Bellcroft credit cards: published only once the company has responded or
    15 days have passed, a few never published, and the CFPB's own range (from January 2025, before the case system)."""
    fh, w = _writer(path, ["Date received", "Product", "Sub-product", "Issue", "Sub-issue", "Company", "State",
                           "Submitted via", "Date sent to company", "Company response to consumer", "Timely response?",
                           "Complaint ID"])
    out = []
    pull = dt.date(2026, 10, 7)
    codes = list(CFPB_ISSUE)
    for r in rows:
        if r["channel"] != "CFPB":
            continue
        sent = r["opened"].date()
        responded = r["events"][-1][1] in ("CLOSED", "MERGED") and r["events"][-1][0].date() <= pull
        if (not responded and (pull - sent).days < 15) or rng.random() < 0.016:
            continue
        iss, sub = CFPB_ISSUE[r["code"]]
        sp = "Store credit card" if "Secured" in r["product"] and rng.random() < 0.1 else "General-purpose credit card or charge card"
        resp = ("Closed with explanation" if rng.random() < 0.78 else "Closed with monetary relief") if responded else "In progress"
        cid = r["source_ref"].split("-")[1]
        out.append([r["cfpb_received"].strftime("%m/%d/%y"), "Credit card", sp, iss, sub, "BELLCROFT BANK, N.A.",
                    STATES[int(rng.integers(0, len(STATES)))], "Web" if rng.random() < 0.88 else "Phone",
                    sent.strftime("%m/%d/%y"), resp, "Yes" if rng.random() < 0.985 else "No", cid])
    d = dt.date(2025, 1, 2)
    while d < dt.date(2025, 9, 26):
        for _ in range(int(rng.poisson(8.4 if d.weekday() < 5 else 2.0))):
            code = codes[int(rng.integers(0, len(codes)))]
            iss, sub = CFPB_ISSUE[code]
            sent = d + dt.timedelta(days=int(rng.integers(1, 6)))
            out.append([d.strftime("%m/%d/%y"), "Credit card", "General-purpose credit card or charge card", iss, sub,
                        "BELLCROFT BANK, N.A.", STATES[int(rng.integers(0, len(STATES)))],
                        "Web" if rng.random() < 0.88 else "Phone", sent.strftime("%m/%d/%y"),
                        "Closed with explanation" if rng.random() < 0.8 else "Closed with monetary relief",
                        "Yes" if rng.random() < 0.98 else "No", "%d" % int(rng.integers(1000000, 9999999))])
        d += dt.timedelta(days=1)
    for row in sorted(out, key=lambda x: (dt.datetime.strptime(x[0], "%m/%d/%y"), x[-1])):
        w.writerow(row)
    fh.close()
    return len(out)


MANIFEST = [
    ("register", "Matter register: every matter opened 1 October 2025 to 8 October 2026", "Bellcroft complaint case system, Fenwright desk extract", "2026-10-09", "Bellcroft confidential (SOW 6, data clause)"),
    ("history", "Matter status history for the same matters", "Bellcroft complaint case system, Fenwright desk extract", "2026-10-09", "Bellcroft confidential (SOW 6, data clause)"),
    ("social", "Social care tool export: posts the desk opened a matter for", "Fenwright social care desk tool", "2026-10-09", "Bellcroft confidential (SOW 6, data clause)"),
    ("meridel", "Meridel's list of the payments received 8 to 10 May 2026", "Meridel Payment Systems, via Bellcroft card operations", "2026-05-22", "Bellcroft confidential"),
    ("portfolio", "Card portfolio: open accounts at each month end", "Bellcroft card finance", "2026-10-06", "Bellcroft confidential"),
    ("codes", "Bellcroft complaint code list, version 7", "Bellcroft complaint program office", "2025-09-24", "Bellcroft confidential"),
    ("time", "Time records of the Fenwright employees badged to the Bellcroft site, 2026", "Fenwright workforce management system", "2026-10-09", "Fenwright internal"),
    ("marlowby", "Marlowby Staffing weekly timesheet exports for the Fenwright placements", "Marlowby Staffing", "2026-10-03", "Supplier data, internal use"),
    ("treadwell", "Treadwell & Oake hours report, 2026", "Treadwell & Oake", "2026-10-02", "Supplier data, internal use"),
    ("opslog", "Complaint desk operations log, 2026", "Fenwright complaint desk", "2026-10-09", "Fenwright internal"),
    ("ivr", "Complaint line weekly call statistics, 2026", "Bellcroft telephony reporting", "2026-10-05", "Bellcroft confidential"),
    ("cfpb", "CFPB consumer complaint database, Bellcroft Bank credit card complaints received from January 2025", "CFPB Consumer Complaint Database (public)", "2026-10-07", "Public data, no restrictions"),
]
DOC_MANIFEST = [
    ("fenwright_wps3_dedicated_team_sizing.pdf", "Workforce Planning Standard WPS-3, version 3.2", "Fenwright Workforce Planning", "2025-02-26", "Fenwright internal"),
    ("meridel_agreement_schedule7_may2026_incident.pdf", "Schedule 7 to the Bellcroft and Meridel ACH processing agreement", "Bellcroft legal, copy to Fenwright", "2026-06-22", "Bellcroft confidential"),
    ("bellcroft_fenwright_sow_2026_exhibit_b.pdf", "Statement of Work No. 6 (2026), Exhibit B", "Bellcroft and Fenwright", "2025-12-11", "Bellcroft confidential"),
    ("data_dictionary_bellcroft_desk.xlsx", "Data dictionary for the files in this folder", "Fenwright workforce planning", "2026-10-09", "Fenwright internal"),
]


def write_manifest(path, counts):
    fh, w = _writer(path, ["file_name", "description", "provided_by", "as_of", "use_terms", "record_count"])
    for k, c, s, d, t in MANIFEST:
        w.writerow([FILES[k], c, s, d, t, counts.get(k, "")])
    for name, c, s, d, t in DOC_MANIFEST:
        w.writerow([name, c, s, d, t, ""])
    fh.close()
    return len(MANIFEST) + len(DOC_MANIFEST)
