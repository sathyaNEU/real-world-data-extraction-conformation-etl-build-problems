"""Independent verifier for task449: reads only the shipped bundle with DuckDB, on its own code path (no calc.py, no
generator state), recomputes the sustained base, the team, the failure's matters a week, the name-and-city share,
the month-end open matters and the monthly flex hours, and compares them with the golden memo, the golden CSV and
figures.json.

    python3 task449/generator/verify.py [target] [golden]
"""
import json
import math
import os
import sys

import duckdb

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)


def rnd(x):
    return int(math.floor(x / 10.0 + 0.5) * 10)


def main(target=None, golden=None):
    target = target or os.path.join(TASK, "target")
    golden = golden or os.path.join(TASK, "golden")
    q = duckdb.connect()
    t = lambda f: os.path.join(target, f).replace("'", "''")
    q.execute("CREATE TABLE reg AS SELECT * FROM read_csv('%s', all_varchar=true)" % t("complaint_matters_2025_2026.csv"))
    q.execute("""CREATE TABLE soc AS SELECT matter_no, CAST(search.results_returned AS INTEGER) AS k
                 FROM read_json('%s', format='newline_delimited')""" % t("social_care_posts_2025_2026.jsonl"))
    q.execute("""CREATE TABLE aff AS SELECT DISTINCT 'BCF' || right(acct_no, 12) AS acct
                 FROM read_csv('%s', all_varchar=true) WHERE coalesce(posted_date, '') <> ''"""
              % t("meridel_affected_payments_20260508_10.csv"))
    q.execute("""CREATE TABLE w AS SELECT r.*, s.k, problem_code IN ('PAY-01','FEE-01','FEE-02','CRB-01') AS four,
                 account_ref IN (SELECT acct FROM aff) AS o, account_basis = 'NAME_CITY' AS nc
                 FROM reg r LEFT JOIN soc s USING (matter_no)
                 WHERE CAST(opened_at AS TIMESTAMP) BETWEEN TIMESTAMP '2026-08-03 00:00:00'
                   AND TIMESTAMP '2026-09-27 23:59:59' AND status <> 'MERGED'""")
    n_all, ver_aff = q.execute("SELECT count(*), sum(CASE WHEN NOT nc AND four AND o THEN 1 ELSE 0 END) FROM w").fetchone()
    nc_n, one_n, one_aff = q.execute("""SELECT count(*), sum(CASE WHEN k = 1 THEN 1 ELSE 0 END),
                                        sum(CASE WHEN k = 1 AND o THEN 1 ELSE 0 END) FROM w WHERE nc AND four""").fetchone()
    share = one_aff / one_n
    pf = q.execute("""SELECT open_accounts FROM read_csv('%s') WHERE month_end = DATE '2026-04-30'"""
                   % t("bellcroft_card_portfolio_2025_2026.csv")).fetchone()[0]
    b = q.execute("SELECT count(*) FROM aff").fetchone()[0] / pf
    so, s1 = q.execute("SELECT sum(o::INT), sum(1.0 / k) FROM w WHERE nc AND four").fetchone()
    inv = (so - b * (nc_n - s1)) / s1
    incident = (ver_aff + share * nc_n) / 8.0
    base = (n_all - ver_aff - share * nc_n) / 8.0
    team = math.ceil(base / 27.5)
    base_inv = (n_all - ver_aff - inv * nc_n) / 8.0
    # month-end open matters from the status history
    q.execute("CREATE TABLE h AS SELECT matter_no, CAST(event_at AS TIMESTAMP) AS evt, event FROM read_csv('%s', all_varchar=true)"
              % t("matter_status_history_2025_2026.csv"))
    open_counts = {}
    for m, last_day in ((1, 31), (2, 28), (3, 31), (4, 30), (5, 31), (6, 30), (7, 31), (8, 31), (9, 30)):
        ts = "2026-%02d-%02d 23:59:59" % (m, last_day)
        open_counts[m] = q.execute("""SELECT count(*) FROM (SELECT matter_no, arg_max(event, evt) AS ev FROM h
                                      WHERE evt <= TIMESTAMP '%s' GROUP BY matter_no) WHERE ev IN ('OPENED','REOPENED')"""
                                   % ts).fetchone()[0]
    # flex hours: Bellcroft overtime, Marlowby's distinct shifts, Treadwell's clock times
    ot = dict(q.execute("""SELECT month(CAST(work_date AS DATE)), sum(CAST(hours AS DOUBLE))
                           FROM read_csv('%s', all_varchar=true) WHERE pay_code = 'OT' AND program = 'BCF-CMP'
                           GROUP BY ALL""" % t("fenwright_time_bellcroft_site_2026.csv")).fetchall())
    mar = dict(q.execute("""SELECT month(strptime(shift_date, '%%m/%%d/%%Y')), sum(CAST(hours AS DOUBLE)) FROM
                            (SELECT DISTINCT * FROM read_csv('%s', all_varchar=true)) GROUP BY ALL"""
                         % t("marlowby_timesheets_fenwright_2026.csv")).fetchall())
    import openpyxl
    tw = openpyxl.load_workbook(t("treadwell_oake_hours_2026.xlsx"), read_only=True)["Hours"]
    tre = {}
    for r in tw.iter_rows(min_row=2, values_only=True):
        st, en, brk = r[2], r[3], int(r[4])
        mins = (int(en[:2]) * 60 + int(en[3:])) - (int(st[:2]) * 60 + int(st[3:])) - brk
        tre[r[1].month] = tre.get(r[1].month, 0.0) + mins / 60.0
    flex = {m: ot.get(m, 0.0) + mar.get(m, 0.0) + tre.get(m, 0.0) for m in range(1, 10)}
    # compare with the goldens
    fail = []
    if team != 35 or math.ceil(base_inv / 27.5) != team or rnd(base) != rnd(base_inv):
        fail.append("team or base")
    fig = json.load(open(os.path.join(HERE, "figures.json")))
    if fig["team"] != str(team) or fig["base"] != "{:,}".format(rnd(base)) or \
            fig["incident_week"] != "{:,}".format(rnd(incident)) or fig["nc_share"] != str(int(round(100 * share))):
        fail.append("figures.json main block")
    rows = open(os.path.join(golden, "flex_hours_2026.csv")).read().strip().splitlines()[1:]
    if [int(r.split(",")[1]) for r in rows] != [rnd(flex[m]) for m in range(1, 10)]:
        fail.append("flex csv")
    from docx import Document
    d = Document(os.path.join(golden, "dedicated_analysts_2027.docx"))
    vals = [d.tables[0].rows[i + 1].cells[1].text for i in range(9)]
    if vals != ["{:,}".format(rnd(open_counts[m])) for m in range(1, 10)]:
        fail.append("open matters table")
    body = " ".join(p.text for p in d.paragraphs)
    if "carry %d dedicated analysts" % team not in body:
        fail.append("memo call")
    print("verify: %d new matters, %d confirmed on the failure, %d name-and-city under the four codes (%d one-match, "
          "%d affected, share %.4f; inversion %.4f); base %.2f, team %d" % (n_all, ver_aff, nc_n, one_n, one_aff, share,
                                                                              inv, base, team))
    print("verify: open %s; flex %s" % (" ".join(str(open_counts[m]) for m in range(1, 10)),
                                        " ".join("%.2f" % flex[m] for m in range(1, 10))))
    if fail:
        raise SystemExit("VERIFY FAILED: " + "; ".join(fail))
    print("verify: PASS (bundle-only recomputation agrees with the memo, the CSV and figures.json)")


if __name__ == "__main__":
    main(*(sys.argv[1:3]))
