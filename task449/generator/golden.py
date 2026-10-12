"""The golden deliverables, computed from the shipped files only: dedicated_analysts_2027.docx (the memo to Bellcroft's
complaint steering committee, with the month-end open matters chart) and flex_hours_2026.csv; and figures.json, every
figure the write-up states."""
import calendar
import csv
import io
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import calc as C     # noqa: E402
import world as W    # noqa: E402

DOCX = "dedicated_analysts_2027.docx"
CSV = "flex_hours_2026.csv"
MONTH = {m: calendar.month_name[m] for m in range(1, 13)}


def fmt(n):
    return "{:,}".format(int(n))


def compute(target):
    B = C.load(target)
    w = C.frame(B)
    s = C.shares(B, w)
    gold = C.base_of(B, "one_match", w)
    rungs = {k: C.base_of(B, k, w) for k in ("attached", "all", "dm_verified", "all_verified", "inversion",
                                             "posterior", "regression", "one_or_two")}
    nc = w[w.nc & w.four]
    one = nc[nc.k == 1]
    X = dict(B=B, w=w, s=s, gold=gold, rungs=rungs, n_one=len(one), n_one_aff=int(one.o.sum()),
             open=C.open_at_month_ends(B, "history"), open_naive=C.open_at_month_ends(B, "register"),
             flex=C.flex_hours(B), flex_naive=C.flex_hours(B, treadwell="decimal"),
             b=C.portfolio_rate(B))
    return X


def figures(X):
    g, s, r = X["gold"], X["s"], X["rungs"]
    return dict(
        team=str(g["team"]),
        base=fmt(C.r10(g["base"])),
        incident_week=fmt(C.r10(g["incident"])),
        nc_share=str(int(round(100 * s["one_match"]))),
        new_matters=fmt(g["N"]),
        new_matters_week=fmt(C.r10(g["N"] / 8.0)),
        verified_affected=fmt(g["ver_aff"]),
        name_city_four=fmt(g["nc_n"]),
        one_match=fmt(X["n_one"]),
        one_match_affected=fmt(X["n_one_aff"]),
        name_city_affected=fmt(C.r10(g["nc_aff"])),
        attached_affected=fmt(round(r["attached"]["nc_aff"])),
        attached_share=str(int(round(100 * s["attached"]))),
        portfolio_rate=str(int(round(100 * X["b"]))),
        default_team=str(r["attached"]["team"]),
        default_base=fmt(C.r10(r["attached"]["base"])),
        dm_team=str(r["dm_verified"]["team"]),
        dm_share=str(int(round(100 * s["dm_verified"]))),
        all_verified_team=str(r["all_verified"]["team"]),
        all_counted_team=str(r["all"]["team"]),
        open_month_end={MONTH[m]: fmt(C.r10(v)) for m, v in X["open"].items()},
        flex={MONTH[m]: fmt(C.r10(v)) for m, v in X["flex"].items()},
        rate="27.5",
        ratio="%.1f" % (g["base"] / C.RATE),
        team_2026=str(W.TEAM_2026),
        clauses=["2.2", "3.1", "4.2", "7.2"],
    )


def write_csv(X, path):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh, lineterminator="\n")
        wr.writerow(["month", "flex_hours"])
        for m in C.MONTHS:
            wr.writerow(["2026-%02d" % m, C.r10(X["flex"][m])])


def chart_png(X):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.ticker import FuncFormatter
    months = [MONTH[m][:3] for m in C.MONTHS]
    vals = [C.r10(X["open"][m]) for m in C.MONTHS]
    fig, ax = plt.subplots(figsize=(6.4, 3.1), dpi=150)
    colours = ["#9aa9b0"] * 4 + ["#8a3b3b"] * 3 + ["#1e5161"] * 2
    bars = ax.bar(months, vals, color=colours, width=0.62)
    for b_, v in zip(bars, vals):
        ax.text(b_.get_x() + b_.get_width() / 2, v + 110, fmt(v), ha="center", va="bottom", fontsize=7.5,
                color="#1f1f1f")
    ax.axvline(3.5, color="#5c5c5c", linewidth=0.8, linestyle=(0, (3, 2)))
    ax.text(3.42, max(vals) * 0.62, "Meridel posting\nfailure, 8 May", fontsize=7.2, color="#5c5c5c", va="center",
            ha="right")
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: "{:,}".format(int(v))))
    ax.set_ylabel("Matters open at month end", fontsize=8.2)
    ax.set_title("Open matters peaked at %s at the end of May and were %s at the end of September"
                 % (fmt(vals[4]), fmt(vals[8])), fontsize=9.2, loc="left")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(labelsize=8)
    ax.set_ylim(0, max(vals) * 1.12)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def _page_field(paragraph):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def write_docx(X, F, path):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.9)
        sec.top_margin = sec.bottom_margin = Inches(0.8)
    h = doc.add_paragraph()
    run = h.add_run("Fenwright Contact Services  |  Workforce Planning")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x5C, 0x5C, 0x5C)
    foot = doc.sections[0].footer.paragraphs[0]
    fr = foot.add_run("Fenwright Contact Services  |  Bellcroft complaint desk: 2027 dedicated team  |  Page ")
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(0x5C, 0x5C, 0x5C)
    _page_field(foot)
    t = doc.add_paragraph()
    rt = t.add_run("%s dedicated analysts carry the Bellcroft desk's sustained workload into 2027" % F["team"])
    rt.bold = True
    rt.font.size = Pt(14)
    rt.font.color.rgb = RGBColor(0x1E, 0x51, 0x61)
    meta = doc.add_paragraph()
    meta.add_run("To: Bellcroft complaint steering committee, meeting of 15 October 2026\n"
                 "From: %s, Workforce Planning, Fenwright Contact Services\nDate: 12 October 2026\n"
                 "Subject: dedicated team in the 2027 statement of work (for approval with the program budget)"
                 % W.PEOPLE["planning"]).font.size = Pt(9.5)
    p = doc.add_paragraph()
    r1 = p.add_run("Recommendation: carry %s dedicated analysts in the 2027 statement of work" % F["team"])
    r1.bold = True
    p.add_run(" (%s in 2026). Under WPS-3 the team is the sustained base, %s new matters a week, over %s matters per "
              "analyst-week, rounded up." % (F["team_2026"], F["base"], F["rate"]))
    doc.add_paragraph("The base. The desk opened %s new matters in the eight complete weeks before October (3 August "
                      "to 27 September), %s a week. About %s a week of them arise from Meridel's May payment-posting "
                      "failure as Schedule 7 defines a covered matter (one of the four codes, from a cardholder whose "
                      "account had an affected payment); the rest, %s a week, is the sustained base."
                      % (F["new_matters"], F["new_matters_week"], F["incident_week"], F["base"]))
    doc.add_paragraph("Counting the failure's matters. For matters with a confirmed account, %s carry an affected "
                      "payment. The social care desk's name-and-city matters cannot be checked one by one: their "
                      "account is the first cardholder the name-and-city search returned, which is the poster's own "
                      "only when the search returned one cardholder. Of those %s matters, %s (%s per cent) carry an "
                      "affected payment, and the same share holds across all %s name-and-city matters under the four "
                      "codes when each attachment is weighted by its chance of being the poster's (the other "
                      "cardholders carry the portfolio's %s per cent). That puts about %s of them on the failure, "
                      "against %s whose attached account happens to be on Meridel's list; counting only those would "
                      "have sized the team at %s."
                      % (F["verified_affected"], F["one_match"], F["one_match_affected"], F["nc_share"],
                         F["name_city_four"], F["portfolio_rate"], F["name_city_affected"], F["attached_affected"],
                         F["default_team"]))
    doc.add_paragraph("Open matters. The backlog peaked at the end of May and has come back toward its winter level; "
                      "the counts below follow each matter's status history, so a matter closed and later reopened "
                      "is not counted as open between the two.")
    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = "Light List Accent 1"
    tbl.rows[0].cells[0].text = "Month end (2026)"
    tbl.rows[0].cells[1].text = "Matters open"
    for m in C.MONTHS:
        row = tbl.add_row().cells
        row[0].text = MONTH[m]
        row[1].text = F["open_month_end"][MONTH[m]]
    src = doc.add_paragraph()
    rs = src.add_run("Source: case system status history, extract of 9 October 2026. Counted at the close of each "
                     "month's last day.")
    rs.italic = True
    rs.font.size = Pt(8)
    doc.add_picture(io.BytesIO(chart_png(X)), width=Inches(6.2))
    doc.add_paragraph("Flex line. Overtime and agency hours by month are in flex_hours_2026.csv; they peaked at %s "
                      "hours in June and were %s hours in September."
                      % (F["flex"]["June"], F["flex"]["September"]))
    notes = doc.add_paragraph()
    rn = notes.add_run("Notes. Weeks run Monday to Sunday, Eastern time, and the eight are the last complete weeks "
                       "before October, when the committee sets the budget (WPS-3 3.1). A matter merged into another "
                       "as a duplicate, or reopened, is not a new matter (WPS-3 2.2). Covered matters are as Schedule 7 "
                       "section 7.2 defines them, whether or not Bellcroft can claim them under section 7.4. Matters "
                       "a week and open matters are rounded to the nearest 10 from unrounded counts; the team is "
                       "rounded up from the unrounded base.")
    rn.font.size = Pt(8.5)
    rn.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    doc.save(path)


def main(target, golden, figures_path):
    os.makedirs(golden, exist_ok=True)
    X = compute(target)
    F = figures(X)
    write_csv(X, os.path.join(golden, CSV))
    write_docx(X, F, os.path.join(golden, DOCX))
    import postprocess as PP
    PP.normalize_ooxml(os.path.join(golden, DOCX), W.PEOPLE["planning"], W.PEOPLE["planning"], "2026-10-12T13:05:00Z",
                       "2026-10-12T16:48:31Z", company=W.BPO)
    stamp = 1791823800.0
    for f in os.listdir(golden):
        os.utime(os.path.join(golden, f), (stamp, stamp))
    with open(figures_path, "w") as fh:
        json.dump(F, fh, indent=1, ensure_ascii=False)
        fh.write("\n")
    return X
