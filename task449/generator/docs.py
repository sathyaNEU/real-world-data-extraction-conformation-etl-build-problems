"""The pack's documents: Fenwright's workforce planning standard WPS-3, Schedule 7 to Bellcroft's processing agreement
with Meridel, Exhibit B of the 2026 statement of work (PDFs), the field dictionary and the desk's operations log
(XLSX)."""
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                           # noqa: E402
from reportlab.lib.pagesizes import LETTER                                 # noqa: E402
from reportlab.lib.styles import ParagraphStyle                            # noqa: E402
from reportlab.lib.units import inch                                       # noqa: E402
from reportlab.pdfbase import pdfmetrics                                   # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                               # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import pack as P    # noqa: E402
import world as W   # noqa: E402

WPS = "fenwright_wps3_dedicated_team_sizing.pdf"
SCHEDULE = "meridel_agreement_schedule7_may2026_incident.pdf"
SOW = "bellcroft_fenwright_sow_2026_exhibit_b.pdf"
DICTIONARY = "data_dictionary_bellcroft_desk.xlsx"

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("LiberationSans", "LiberationSans-Regular.ttf"), ("LiberationSans-Bold", "LiberationSans-Bold.ttf"),
              ("LiberationSans-Italic", "LiberationSans-Italic.ttf"), ("LiberationSerif", "LiberationSerif-Regular.ttf"),
              ("LiberationSerif-Bold", "LiberationSerif-Bold.ttf"), ("LiberationSerif-Italic", "LiberationSerif-Italic.ttf")):
    pdfmetrics.registerFont(TTFont(nm, FONT_DIR + f))
INK = colors.HexColor("#1f1f1f")
GREY = colors.HexColor("#5c5c5c")
TEAL = colors.HexColor("#1e5161")
MAROON = colors.HexColor("#5a2232")

F_BODY = ParagraphStyle("fb", fontName="LiberationSans", fontSize=9.6, leading=13.2, textColor=INK, spaceAfter=5)
F_NUM = ParagraphStyle("fn", parent=F_BODY, leftIndent=24, firstLineIndent=-24)
F_H1 = ParagraphStyle("fh1", fontName="LiberationSans-Bold", fontSize=14.5, leading=18, textColor=TEAL, spaceAfter=2)
F_SUB = ParagraphStyle("fs", fontName="LiberationSans", fontSize=9.2, leading=12, textColor=GREY, spaceAfter=10)
F_H2 = ParagraphStyle("fh2", fontName="LiberationSans-Bold", fontSize=10.6, leading=14, textColor=TEAL, spaceBefore=7,
                      spaceAfter=3)
L_BODY = ParagraphStyle("lb", fontName="LiberationSerif", fontSize=10.4, leading=14.0, textColor=INK, spaceAfter=5)
L_NUM = ParagraphStyle("ln", parent=L_BODY, leftIndent=28, firstLineIndent=-28)
L_H1 = ParagraphStyle("lh1", fontName="LiberationSerif-Bold", fontSize=13.5, leading=17, textColor=MAROON, spaceAfter=3,
                      alignment=1)
L_SUB = ParagraphStyle("ls", fontName="LiberationSerif-Italic", fontSize=10, leading=13, textColor=GREY, spaceAfter=10,
                       alignment=1)
L_H2 = ParagraphStyle("lh2", fontName="LiberationSerif-Bold", fontSize=11, leading=14.5, textColor=MAROON,
                      spaceBefore=7, spaceAfter=3)


def _frame(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("LiberationSans", 7.4)
        canvas.setFillColor(GREY)
        canvas.drawString(0.85 * inch, 0.55 * inch, text)
        canvas.drawRightString(LETTER[0] - 0.85 * inch, 0.55 * inch, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _doc(path, title, author):
    return SimpleDocTemplate(path, pagesize=LETTER, leftMargin=0.9 * inch, rightMargin=0.9 * inch,
                             topMargin=0.85 * inch, bottomMargin=0.9 * inch, title=title, author=author, subject="",
                             creator="", producer="")


def write_wps(path):
    s = [Paragraph("Workforce Planning Standard WPS-3", F_H1),
         Paragraph("Dedicated team sizing for client programs. Version 3.2, effective 1 March 2025. Owner: Workforce "
                   "Planning, %s." % W.PEOPLE["planning"], F_SUB)]
    N = lambda n, t: s.append(Paragraph("<b>%s</b>&nbsp;&nbsp;%s" % (n, t), F_NUM))
    s.append(Paragraph("1. Scope", F_H2))
    N("1.1", "This standard sets how Fenwright sizes the dedicated team that a client program's statement of work "
             "carries into the next contract year. It applies to complaint and dispute programs billed per dedicated "
             "analyst-month with a flex line.")
    s.append(Paragraph("2. Definitions", F_H2))
    N("2.1", "<i>Matter</i>: one customer's complaint about one problem with one account, as the client's case "
             "system records it.")
    N("2.2", "<i>New matter</i>: a matter opened for the first time in the period. A matter reopened after it was "
             "closed keeps its first opening and is not a new matter when it is reopened. A matter merged into another "
             "as a duplicate is not a new matter.")
    N("2.3", "<i>Processor or vendor incident</i>: a failure at a third party the client uses, where the client's "
             "agreement with that party defines the matters arising from it.")
    s.append(Paragraph("3. The sustained base", F_H2))
    N("3.1", "The sustained base is the mean of the new matters in each of the last eight complete weeks before the "
             "month in which the client sets the budget, less the new matters in those weeks arising from a processor "
             "or vendor incident, as the incident's agreement defines the matters it covers. Weeks run Monday to "
             "Sunday in the time zone of the client's case system.")
    N("3.2", "Matters arising from an incident are left out because the dedicated team is sized for the workload "
             "that recurs. They are worked on the flex line while they last.")
    s.append(Paragraph("4. The dedicated team", F_H2))
    N("4.1", "The planning rate is %.1f new matters per analyst-week, across channels." % W.RATE)
    N("4.2", "The dedicated team is the sustained base divided by the planning rate, rounded up to a whole analyst.")
    s.append(Paragraph("5. The flex line", F_H2))
    N("5.1", "Overtime and agency hours carry the workload above the dedicated team's capacity, incident matters "
             "included. The flex line is budgeted with the team but sized separately.")
    s.append(Paragraph("6. Review", F_H2))
    N("6.1", "Workforce Planning reviews the planning rate each year against the matters the programs handled. "
             "Changes take effect from the next budget round.")
    d = _doc(path, "WPS-3 Dedicated team sizing", "Fenwright Contact Services")
    fr = _frame("Fenwright Contact Services  |  WPS-3 Dedicated team sizing  |  v3.2")
    d.build(s, onFirstPage=fr, onLaterPages=fr)


def write_schedule(path):
    s = [Paragraph("Amendment No. 3 to the ACH Payment Processing Agreement", L_H1),
         Paragraph("between Bellcroft Bank, N.A. and Meridel Payment Systems, Inc.<br/>Schedule 7: Posting incident of "
                   "May 2026. Signed 19 June 2026.", L_SUB)]
    N = lambda n, t: s.append(Paragraph("<b>%s</b>&nbsp;&nbsp;%s" % (n, t), L_NUM))
    s.append(Paragraph("Schedule 7", L_H2))
    N("7.1", "<i>Affected Payment</i>: an ACH payment to a Bellcroft card account that Meridel received on 8, 9 or 10 "
             "May 2026 and that posted to the account after the business day following its receipt. A payment "
             "returned unpaid did not post and is not an Affected Payment; a returned payment presented again and "
             "posted late is. Meridel delivered its list of these payments to Bellcroft on 22 May 2026.")
    N("7.2", "<i>Covered Matter</i>: a complaint that Bellcroft, or the complaint desk Bellcroft engages, opens from "
             "8 May 2026 to 31 December 2026 under complaint code PAY-01, FEE-01, FEE-02 or CRB-01 of Bellcroft's "
             "complaint code list, made by or for a cardholder whose account had an Affected Payment.")
    N("7.3", "<i>Remediation</i>: Meridel funds Bellcroft's reversal of the late fees and interest charged on accounts "
             "with an Affected Payment, and Bellcroft's requests to the credit bureaus to remove late-payment "
             "reporting an Affected Payment caused.")
    N("7.4", "<i>Handling cost</i>: Meridel pays Bellcroft $41.00 for each Covered Matter Bellcroft claims. Bellcroft "
             "claims monthly, listing each Covered Matter with the account number the cardholder confirmed. Meridel "
             "pays within 30 days of a claim.")
    N("7.5", "This Schedule ends on 31 March 2027. Sections 14 (Audit) and 19 (Confidentiality) of the Agreement apply "
             "to it.")
    s.append(Spacer(1, 14))
    sig = Table([["For Bellcroft Bank, N.A.", "For Meridel Payment Systems, Inc."],
                 ["Head of Card Operations", "Chief Operating Officer"], ["19 June 2026", "19 June 2026"]],
                colWidths=[3.2 * inch, 3.2 * inch])
    sig.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "LiberationSerif", 9.6), ("TEXTCOLOR", (0, 0), (-1, -1), INK),
                             ("LINEABOVE", (0, 1), (-1, 1), 0.5, GREY), ("TOPPADDING", (0, 0), (-1, -1), 4)]))
    s.append(sig)
    d = _doc(path, "Schedule 7 Posting incident of May 2026", "Bellcroft Bank, N.A.")
    fr = _frame("Bellcroft / Meridel ACH Payment Processing Agreement  |  Amendment No. 3  |  Schedule 7")
    d.build(s, onFirstPage=fr, onLaterPages=fr)


def write_sow(path):
    s = [Paragraph("Statement of Work No. 6 (2026)", L_H1),
         Paragraph("under the Master Services Agreement between Bellcroft Bank, N.A. and Fenwright Contact Services, LLC"
                   "<br/>Exhibit B: Staffing and charges", L_SUB)]
    N = lambda n, t: s.append(Paragraph("<b>%s</b>&nbsp;&nbsp;%s" % (n, t), L_NUM))
    N("B.1", "<i>Services</i>: Fenwright operates Bellcroft's card complaint desk from its Columbus site. The desk "
             "opens and works complaints received by phone, through Bellcroft's digital channels (secure message and "
             "the online complaint form), from the CFPB, from state regulators, through Bellcroft's executive office "
             "and on social media.")
    N("B.2", "<i>Dedicated team</i>: %d dedicated analysts for 2026, billed at $6,940 per analyst-month. Dedicated "
             "analysts work the Bellcroft program only." % W.TEAM_2026)
    N("B.3", "<i>Flex hours</i>: hours worked on the Bellcroft program beyond the dedicated team's scheduled hours, "
             "that is, overtime worked on the program by Fenwright analysts and hours worked by agency analysts placed "
             "on it. Flex hours are billed at $58.00 an hour in the month they are worked.")
    N("B.4", "<i>Hours of operation</i> (Eastern time): phone lines 8 am to 9 pm Monday to Friday and 9 am to 5 pm "
             "Saturday; the social care desk 7 am to 11 pm daily. Digital, CFPB, state and executive office complaints "
             "are opened on receipt or on the next business morning.")
    N("B.5", "<i>Next year's team</i>: each October Bellcroft and Fenwright set the next contract year's dedicated "
             "team with Bellcroft's complaint program budget, under Fenwright's workforce planning standard WPS-3. "
             "Bellcroft's complaint steering committee approves it.")
    N("B.6", "<i>Reporting</i>: Fenwright reports new, open and closed matters weekly, and flex hours and analyst "
             "attendance monthly, to Bellcroft's complaint program manager.")
    d = _doc(path, "SOW 6 (2026) Exhibit B", "Fenwright Contact Services")
    fr = _frame("Bellcroft / Fenwright  |  SOW No. 6 (2026)  |  Exhibit B")
    d.build(s, onFirstPage=fr, onLaterPages=fr)


FIELDS = [
    ("register", "matter_no", "Case system matter number"),
    ("register", "opened_at", "Date and time the case system opened the matter, Eastern time"),
    ("register", "channel", "Phone, Digital (secure message or online form), CFPB, State (state regulator referral), "
                            "Executive (executive office) or Social (social care desk)"),
    ("register", "source_ref", "The channel's own reference for the contact that opened the matter: IVR call id, "
                               "message id, CFPB complaint id, state referral number, executive office letter number "
                               "or social post id"),
    ("register", "problem_code", "Complaint code (complaint_codes_v7.csv)"),
    ("register", "product", "Card product on the account"),
    ("register", "account_basis", "How the account on the matter was identified: VERIFIED, confirmed with the "
                                  "cardholder (by the phone system, a signed-in session, the regulator's or "
                                  "executive office's details, or a direct message); NAME_CITY, attached by the social "
                                  "care desk from a search of cardholders by the poster's name and city"),
    ("register", "account_ref", "Bellcroft account reference of the account on the matter: BCF and the twelve-digit "
                                "account number. For NAME_CITY, the first cardholder the search returned, in "
                                "account-reference order"),
    ("register", "queue", "Desk queue: GEN, REG (regulatory) or SOC (social care)"),
    ("register", "status", "OPEN, CLOSED or MERGED (closed as a duplicate of another matter)"),
    ("register", "closed_at", "Date and time the matter was closed, Eastern time"),
    ("history", "matter_no", "Case system matter number"),
    ("history", "event_at", "Date and time of the status change, Eastern time"),
    ("history", "event", "OPENED, CLOSED, REOPENED or MERGED"),
    ("social", "post_id", "Platform's id for the post"),
    ("social", "created_at", "Time the post was published, UTC"),
    ("social", "handle, display_name, city, state", "The poster's handle and display name, and the city and state "
                                                     "the poster gave"),
    ("social", "search.name, search.city", "The name and city the desk searched cardholders by"),
    ("social", "search.results_returned", "Number of cardholders the search returned"),
    ("social", "account.basis, account.ref", "The account on the matter and how it was identified (as in the "
                                             "matter register)"),
    ("social", "dm.requested, dm.verified", "The desk asked the poster to send a direct message; the poster "
                                            "confirmed their account by direct message"),
    ("social", "matter_no", "Matter the desk opened for the post"),
    ("meridel", "payment_ref", "Meridel's payment reference"),
    ("meridel", "acct_no", "Bellcroft account number, zero-padded to fourteen digits"),
    ("meridel", "received_date", "Date Meridel received the payment"),
    ("meridel", "due_post_date", "Date the payment should have posted"),
    ("meridel", "posted_date", "Date the payment posted to the account"),
    ("meridel", "business_days_late", "Business days between due_post_date and posted_date"),
    ("meridel", "status", "POSTED_LATE; RETURNED (returned unpaid); REPRESENTED (a returned payment presented again)"),
    ("meridel", "return_code", "ACH return reason code"),
    ("meridel", "original_payment_ref", "For REPRESENTED, the returned payment"),
    ("portfolio", "month_end", "Month end"),
    ("portfolio", "open_accounts", "Open card accounts at month end"),
    ("portfolio", "new_accounts, closed_accounts", "Accounts opened and closed in the month"),
    ("codes", "code, description, family", "Bellcroft complaint code list, version 7"),
    ("time", "employee_id", "Fenwright employee id"),
    ("time", "work_date, segment_start, segment_end", "Date and clock times of the time segment"),
    ("time", "pay_code", "REG regular, OT overtime, PTO paid time off, HOL holiday"),
    ("time", "program", "Client program: BCF-CMP (Bellcroft complaints) or WXE-CS (Wrexley Energy customer service)"),
    ("time", "hours", "Hours in the segment"),
    ("marlowby", "worker_no", "Marlowby Staffing worker number"),
    ("marlowby", "shift_date, time_in, time_out", "Date and clock times of the shift"),
    ("marlowby", "unpaid_break_min", "Unpaid break, minutes"),
    ("marlowby", "hours", "Hours worked"),
    ("marlowby", "week_ending, export_id", "Billing week (ending Saturday) and the weekly export the row came in"),
    ("treadwell", "Associate ID", "Treadwell & Oake associate"),
    ("treadwell", "Work Date, Start, End", "Date and clock times of the shift"),
    ("treadwell", "Break (min)", "Unpaid break, minutes"),
    ("treadwell", "Hours", "Hours worked"),
    ("treadwell", "Approved By", "Fenwright supervisor who approved the shift"),
    ("opslog", "Date, By, Entry", "Date, author and text of the desk's log entry"),
    ("ivr", "week_starting", "Monday of the week"),
    ("ivr", "calls_offered, calls_answered, calls_abandoned", "Calls to the complaint line in the week"),
    ("ivr", "ASA (sec), AHT (sec)", "Average speed of answer and average handle time, seconds"),
    ("cfpb", "(all)", "The CFPB's published fields for complaints about Bellcroft credit cards; Complaint ID is the CFPB "
                      "complaint id (source_ref of CFPB matters without the date prefix)"),
]
GRAIN = [
    ("register", "A matter the case system opened from its first day, 1 October 2025, to 8 October 2026. "
                 "Complaints opened earlier were worked to closure in the previous case system."),
    ("history", "A change in a matter's status."),
    ("social", "A post the social care desk opened a matter for (one JSON record)."),
    ("meridel", "A payment on Meridel's list for 8 to 10 May 2026."),
    ("portfolio", "A month end."),
    ("time", "A time segment of a Fenwright employee badged to the Bellcroft site, 2026."),
    ("marlowby", "A shift, as it appeared in a weekly export Marlowby Staffing sent."),
    ("treadwell", "A shift."),
    ("opslog", "A log entry."),
    ("ivr", "A week."),
    ("cfpb", "A complaint the CFPB published, by the CFPB's own dates."),
]


def write_dictionary(path):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    name = dict(P.FILES)
    owner = {"register": "Bellcroft case system", "history": "Bellcroft case system", "social": "Social care desk tool",
             "meridel": "Meridel Payment Systems", "portfolio": "Bellcroft card finance", "codes": "Bellcroft complaint program",
             "time": "Fenwright workforce management", "marlowby": "Marlowby Staffing", "treadwell": "Treadwell & Oake",
             "opslog": "Fenwright complaint desk", "ivr": "Bellcroft telephony", "cfpb": "CFPB"}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Files"
    ws.append(["File", "One row is", "Kept by"])
    for f, g in GRAIN:
        ws.append([name[f], g, owner[f]])
    ws2 = wb.create_sheet("Columns")
    ws2.append(["Source file", "Column", "Description"])
    for f, fld, d in FIELDS:
        ws2.append([name[f], fld, d])
    for sh, widths in ((ws, (40, 100, 28)), (ws2, (40, 30, 100))):
        for c in sh[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1E5161")
        for col, wd in zip("ABC", widths):
            sh.column_dimensions[col].width = wd
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        sh.freeze_panes = "A2"
    wb.save(path)


OPSLOG = [
    ("2026-01-02", "Desk manager", "Holiday rota closed out; normal hours from today."),
    ("2026-01-05", "CS", "Treadwell & Oake evening and weekend cover renewed for 2026 (five associates, approvals by "
                         "the evening supervisors)."),
    ("2026-01-13", "Desk manager", "Case system release 4.2 overnight: queue filters and the letter templates changed; "
                                   "no data changes."),
    ("2026-01-27", "Desk manager", "IVR outage 10:05 to 10:40; callers advised to call back. Bellcroft telephony "
                                   "ticket raised."),
    ("2026-02-09", "Desk manager", "Bellcroft complaint program office walk-through of the desk, 10:00 to 12:00."),
    ("2026-02-16", "Desk manager", "Wrexley Energy billing changeover: the five cross-trained analysts work the "
                                   "Wrexley queue until further notice."),
    ("2026-03-02", "Desk manager", "Quarterly calibration on complaint codes with Bellcroft's complaint program "
                                   "office; no code changes."),
    ("2026-03-23", "CS", "Evening supervisor rota changed: Treadwell & Oake approvals by the two evening leads."),
    ("2026-04-14", "Desk manager", "Fire drill 14:00; phone queue paused 12 minutes."),
    ("2026-05-08", "Desk manager", "Meridel ACH posting failure: cardholder payments received 8 to 10 May are not "
                                   "posting. Bellcroft holding statement on the IVR and social channels."),
    ("2026-05-11", "CS", "Overtime authorised for all dedicated analysts until further notice."),
    ("2026-05-12", "CS", "Marlowby Staffing placements start on the Bellcroft queues; Treadwell & Oake adding two "
                         "associates for May and June."),
    ("2026-05-18", "SF", "Social care desk: matters opened on complaint posts from today, not only on direct messages."),
    ("2026-05-22", "Desk manager", "Meridel's list of the 8 to 10 May payments received from Bellcroft card operations."),
    ("2026-06-19", "Desk manager", "Bellcroft and Meridel signed Schedule 7 (handling reimbursement and remediation)."),
    ("2026-06-29", "Desk manager", "Case system release 4.3: bulk close for duplicate matters limited to team leads."),
    ("2026-07-03", "Desk manager", "Independence Day observed: phone lines closed; social care desk open."),
    ("2026-07-20", "Desk manager", "Bellcroft remediation letters mailing 20 to 24 July."),
    ("2026-08-17", "CS", "Two Treadwell & Oake associates on Saturday cover for August."),
    ("2026-08-03", "CS", "Overtime authorisation narrowed to the regulatory and social care queues."),
    ("2026-09-07", "Desk manager", "Labor Day: phone lines closed; social care desk open."),
    ("2026-09-15", "Desk manager", "Annual complaint handling refresher for all analysts, in two sessions."),
    ("2026-09-30", "CS", "Marlowby Staffing placements end."),
    ("2026-10-09", "CS", "Workforce planning pull for the 2027 budget; Bellcroft's complaint steering committee meets "
                         "15 October."),
]


def write_opslog(path):
    import datetime as dt
    import openpyxl
    from openpyxl.styles import Alignment, Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Log"
    ws.append(["Date", "By", "Entry"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for d, by, e in sorted(OPSLOG, key=lambda x: x[0]):
        ws.append([dt.date.fromisoformat(d), by, e])
    for row in ws.iter_rows(min_row=2, max_col=1):
        for c in row:
            c.number_format = "d mmm yyyy"
    for col, wd in zip("ABC", (12, 14, 110)):
        ws.column_dimensions[col].width = wd
    for row in ws.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
