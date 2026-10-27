"""The pack's documents: the phase 3 dose-selection charter, the SAP's definitions section, section 5 of the
investigator's brochure and the phase 1 pharmacokinetic summary tables."""
from __future__ import annotations

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import LETTER                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import inch                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import common as K    # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/"
for _n, _f in (("Carlito", "crosextra/Carlito-Regular.ttf"), ("Carlito-Bold", "crosextra/Carlito-Bold.ttf"),
               ("Carlito-Italic", "crosextra/Carlito-Italic.ttf")):
    pdfmetrics.registerFont(TTFont(_n, FONT_DIR + _f))
SANS, SANS_B, SANS_I = "Carlito", "Carlito-Bold", "Carlito-Italic"
INK = colors.HexColor("#1d2330")
GREY = colors.HexColor("#5a6270")
RULE = colors.HexColor("#a3abb6")
PALE = colors.HexColor("#edf0f4")
TEAL = colors.HexColor("#14505c")

BODY = ParagraphStyle("body", fontName=SANS, fontSize=10.4, leading=13.8, textColor=INK, spaceAfter=6)
CL = ParagraphStyle("cl", parent=BODY, leftIndent=30, firstLineIndent=-30)
H1 = ParagraphStyle("h1", fontName=SANS_B, fontSize=15, leading=19, textColor=TEAL, spaceAfter=2)
H1S = ParagraphStyle("h1s", fontName=SANS, fontSize=9.6, leading=12.4, textColor=GREY, spaceAfter=10)
H2 = ParagraphStyle("h2", fontName=SANS_B, fontSize=11.2, leading=14, textColor=INK, spaceBefore=8, spaceAfter=3)
NOTE = ParagraphStyle("note", fontName=SANS_I, fontSize=9, leading=11.6, textColor=GREY, spaceBefore=6)
CELL = ParagraphStyle("cell", fontName=SANS, fontSize=9, leading=11, textColor=INK)
CELLB = ParagraphStyle("cellb", fontName=SANS_B, fontSize=9, leading=11, textColor=INK)


def _footer(left, right):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont(SANS, 7.6)
        canvas.setFillColor(GREY)
        canvas.drawString(0.9 * inch, 0.55 * inch, left)
        canvas.drawRightString(LETTER[0] - 0.9 * inch, 0.55 * inch, "%s  |  page %d" % (right, doc.page))
        canvas.restoreState()
    return draw


def _doc(path, title, author):
    return SimpleDocTemplate(path, pagesize=LETTER, leftMargin=1.0 * inch, rightMargin=1.0 * inch, topMargin=0.85 * inch,
                             bottomMargin=0.9 * inch, title=title, author=author, subject="", creator="", producer="")


def _table(rows, widths):
    data = [[Paragraph(str(c), CELLB if i == 0 else CELL) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, hAlign="LEFT")
    t.setStyle(TableStyle([("LINEABOVE", (0, 0), (-1, 0), 0.6, RULE), ("LINEBELOW", (0, 0), (-1, 0), 0.6, RULE),
                           ("LINEBELOW", (0, -1), (-1, -1), 0.6, RULE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("BACKGROUND", (0, 0), (-1, 0), PALE), ("TOPPADDING", (0, 0), (-1, -1), 2.4),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 2.4)]))
    return t


def write_charter(path):
    s = []
    s.append(Paragraph("Dovikizumab (CVN-2184) phase 3 dose-selection charter", H1))
    s.append(Paragraph("Corvane Therapeutics, Clinical Pharmacology and Biostatistics. Version 1.0, adopted 9 June 2026, "
                       "before the CVN2184-UC-201 database lock and unblinding. Document CVN-CP-0142.", H1S))
    paras = [
        ("1", "Purpose", "This charter fixes, before the phase 2b results are unblinded, how Corvane will choose the "
         "induction regimen it proposes for the phase 3 ulcerative colitis programme at the end-of-phase-2 meeting."),
        ("2", "Candidates", "Dovikizumab 300 mg, 600 mg or 900 mg, given intravenously at weeks 0, 4 and 8. No other dose, "
         "interval or route is a candidate."),
        ("3", "What is forecast", "For each candidate, the week-12 clinical remission rate that patients meeting the "
         "CVN2184-UC-201 eligibility criteria would show on that regimen. The phase 3 protocol repeats those criteria "
         "unchanged. Clinical remission is the protocol's modified Mayo definition, scored with the central "
         "adjudicated endoscopy read, and a patient who discontinues study treatment or takes rescue therapy before "
         "week 12 is a non-remitter (non-responder imputation), as in the phase 2b primary analysis."),
        ("4", "Basis", "The forecasts are model-informed: an exposure-response analysis of the CVN2184-UC-201 dovikizumab "
         "arm, with exposure at the other candidates obtained by proportional scaling of dose, which the phase 1 "
         "programme supports (investigator's brochure, section 5). Each forecast is the remission rate the model "
         "predicts over the exposure-response analysis set, the dovikizumab patients for whom the analysis's exposure "
         "measure is available. The model is a logistic regression of week-12 clinical remission on the natural "
         "logarithm of the exposure measure, with the randomisation stratum (prior biologic or JAK inhibitor failure) "
         "as a covariate. The charter does not prescribe the exposure measure; the analysis records its choice and "
         "its reasons."),
        ("5", "Rule", "Phase 3 carries the lowest candidate whose forecast week-12 clinical remission rate is no more "
         "than 5 percentage points below the highest of the three forecasts. Forecasts are compared unrounded."),
        ("6", "Records", "The analysis uses the CVN2184-UC-201 database at its lock, planned for July 2026, the bioanalytical "
         "results Brightwater Bioanalytics reported for it, and the study's other locked datasets."),
        ("7", "Ownership", "The head of clinical pharmacology owns the forecasts and the proposal; the biostatistics lead "
         "reviews them. Safety and immunogenicity summaries for the briefing package follow the statistical analysis "
         "plan's definitions (section 6).")]
    for n, h, t in paras:
        s.append(Paragraph("%s. %s" % (n, h), H2))
        s.append(Paragraph(t, BODY))
    s.append(Paragraph("Approved for Corvane Therapeutics by the Clinical Development Committee, 9 June 2026.", NOTE))
    d = _doc(path, "CVN-2184 phase 3 dose-selection charter", "Corvane Therapeutics")
    d.build(s, onFirstPage=_footer("CVN-CP-0142 v1.0", "Corvane Therapeutics confidential"),
            onLaterPages=_footer("CVN-CP-0142 v1.0", "Corvane Therapeutics confidential"))


def write_sap(path):
    s = []
    s.append(Paragraph("CVN2184-UC-201 statistical analysis plan, section 6: definitions", H1))
    s.append(Paragraph("Excerpt. SAP version 2.0, final, signed 2 April 2026 before the database lock. Sections 1 to 5 "
                       "(objectives, design, sample size, analysis sets, multiplicity) and 7 to 9 are not reproduced.", H1S))
    items = [
        ("6.1", "Study day", "Day 1 is the day of the first infusion. Visits at weeks 4, 8 and 12 fall on days 29, 57 and "
         "85, with a window of three days either side for week 12; infusions at weeks 4 and 8 are given on schedule."),
        ("6.2", "Partial dates", "An adverse event onset date with an unknown day is the first day of that month, unless "
         "the first infusion fell in that month, in which case it is the date of the first infusion. Dates with an "
         "unknown month are not used."),
        ("6.3", "Treatment-emergent adverse event", "An adverse event with onset on or after the first infusion and no "
         "later than the earlier of 30 days after the last infusion and the patient's last study visit. Sites record an "
         "event that began before the first infusion and worsened after it as a new event with its own onset date, so "
         "whether an event is treatment-emergent is decided by its onset date as recorded; an event's seriousness or "
         "severity changing on a later review does not make it a new event. The adverse event listing carries one "
         "record for each event at each review while it is ongoing, each showing the event as it stood at that "
         "review; an event's latest record is its current form."),
        ("6.4", "Serious infection", "A treatment-emergent serious adverse event in the MedDRA system organ class "
         "Infections and infestations."),
        ("6.5", "Infusion-related reaction", "A treatment-emergent adverse event with preferred term Infusion related "
         "reaction, Hypersensitivity, Urticaria, Flushing, Chills or Infusion site reaction, with onset on the day of an "
         "infusion or the day after."),
        ("6.6", "Exposure", "Patient-years from the first infusion to the earlier of 30 days after the last infusion and "
         "the last study visit, both days counted, over 365.25. Exposure-adjusted event rates count events per 100 "
         "patient-years; a patient with two serious infections contributes two events."),
        ("6.7", "Immunogenicity", "Samples are tested in a tiered assay: screening (signal to noise against the "
         "screening cut point), confirmatory (per cent inhibition against the confirmatory cut point) and titer. A "
         "patient's sample at a visit (the blood drawn at it) is ADA positive when it is screening positive and confirmed. A patient is treatment-emergent ADA "
         "positive at a visit when ADA negative at baseline and positive at the visit (treatment-induced), or positive "
         "at both with a titer at the visit at least four times the baseline titer (treatment-boosted). Incidence at a "
         "visit is over patients with an evaluable baseline and an evaluable sample at that visit."),
        ("6.8", "Pharmacokinetic samples", "A pre-dose sample is drawn before the infusion starts; a sample drawn after "
         "the start is not a pre-dose sample and is not used as one. Concentrations below the lower limit of "
         "quantification (0.100 ug/mL) are reported as below it; they are not quantified and are not used as "
         "exposure values. Sample times are in UTC; infusion times are the site's local time."),
        ("6.9", "Clinical remission", "Modified Mayo score of 2 or less with a stool frequency subscore of 1 or less, a "
         "rectal bleeding subscore of 0 and an endoscopic subscore of 1 or less without friability, the endoscopic "
         "subscore from the central adjudicated read; non-responder imputation as in the charter."),
        ("6.10", "Symptomatic remission (e-diary)", "At a reference day, stool frequency and rectal bleeding subscores "
         "from the patient's e-diary: the mean of the three most recent days with an entry among the seven days before "
         "the reference day, each mean rounded to the nearest whole number (halves up); the stool frequency subscore "
         "compares the rounded mean with the patient's normal stool frequency at screening (0 at or below it, 1 for one "
         "or two above, 2 for three or four above, 3 for five or more above). Symptomatic remission is a stool frequency "
         "subscore of 1 or less and a rectal bleeding subscore of 0. A patient's value for a day is the last entry the "
         "patient made for that day. Reference days for weeks 6, 8 and 10 are days 43, 57 and 71; a patient counts at a "
         "reference day while still in the study on it. Patients with fewer than three days with entries in the seven "
         "days are not in symptomatic remission.")]
    for n, h, t in items:
        s.append(Paragraph("%s %s" % (n, h), H2))
        s.append(Paragraph(t, BODY))
    d = _doc(path, "CVN2184-UC-201 SAP section 6", "Corvane Therapeutics Biostatistics")
    d.build(s, onFirstPage=_footer("CVN2184-UC-201 SAP v2.0", "Corvane Therapeutics confidential"),
            onLaterPages=_footer("CVN2184-UC-201 SAP v2.0", "Corvane Therapeutics confidential"))


def write_ib(path):
    s = []
    s.append(Paragraph("Dovikizumab investigator's brochure, edition 6", H1))
    s.append(Paragraph("Section 5: pharmacokinetics and product metabolism in humans (excerpt). Effective 15 January 2026.", H1S))
    paras = [
        "5.1 Dovikizumab is a humanised IgG1 monoclonal antibody against the p19 subunit of interleukin-23. As for other "
        "IgG1 antibodies, its elimination is by proteolytic catabolism; no renal or hepatic route contributes.",
        "5.2 In study CVN2184-HV-101, 48 healthy volunteers received single IV doses of 100, 300, 600 or 1,200 mg. "
        "Maximum concentration and area under the curve rose in proportion to dose across the range (power-model slopes "
        "0.99 and 1.02), with a terminal half-life of 15 to 18 days and clearance of about 0.24 L/day.",
        "5.3 A population model of the phase 1 data found body weight and serum albumin to be the covariates on "
        "clearance (higher weight and lower albumin, higher clearance) and body weight on volume. No dose or time "
        "dependence of clearance was identified in healthy volunteers.",
        "5.4 Anti-drug antibodies were detected in 2 of 48 volunteers, at low titer and without an apparent effect on "
        "exposure."]
    for p in paras:
        s.append(Paragraph(p, BODY))
    d = _doc(path, "Dovikizumab IB edition 6 section 5", "Corvane Therapeutics")
    d.build(s, onFirstPage=_footer("IB edition 6", "Corvane Therapeutics confidential"),
            onLaterPages=_footer("IB edition 6", "Corvane Therapeutics confidential"))


def write_phase1(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HV-101 NCA"
    ws.append(["Cohort", "Dose (mg IV)", "N", "Cmax (ug/mL), geometric mean", "AUCinf (ug.day/mL), geometric mean",
               "t1/2 (days), median", "CL (L/day), geometric mean"])
    for row in (("A", 100, 12, 33.1, 418, 15.6, 0.239), ("B", 300, 12, 101.6, 1262, 16.4, 0.238),
                ("C", 600, 12, 197.3, 2541, 17.1, 0.236), ("D", 1200, 12, 409.8, 5134, 16.9, 0.234)):
        ws.append(list(row))
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True)
    for col, w in zip("ABCDEFG", (9, 12, 6, 18, 22, 14, 16)):
        ws.column_dimensions[col].width = w
    ws.append([])
    ws.append(["Source: CVN2184-HV-101 clinical study report, table 11.2-1 (healthy volunteers)."])
    wb.save(path)
