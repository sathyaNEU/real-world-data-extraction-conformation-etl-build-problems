"""The golden deliverables: the induction dose memo (docx), the phase 2b induction tables (xlsx) and the
exposure-response chart (png), every figure from the cut pack."""
from __future__ import annotations

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import common as K

INK = "#1d2330"
GREY = "#6b7280"
TEAL = "#14505c"
SAND = "#c8873a"


def chart(path, F):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9.5, "axes.edgecolor": "#9aa3ad",
                         "axes.labelcolor": INK, "xtick.color": INK, "ytick.color": INK, "svg.hashsalt": "t171"})
    fig, (a, b) = plt.subplots(1, 2, figsize=(10.4, 4.3), dpi=150, gridspec_kw={"width_ratios": [1.35, 1]})
    q = F['quartiles']
    xs = np.arange(4)
    a.bar(xs, [v['rate'] for v in q], width=0.58, color="#9cc3c9", edgecolor="white", linewidth=2, zorder=2)
    for x, v in zip(xs, q):
        a.text(x, v['rate'] / 2, "%.0f%%" % v['rate'], ha="center", va="center", color="white", fontsize=9, fontweight="bold")
        a.text(x, 1.5, "n=%d" % v['n'], ha="center", va="bottom", color="white", fontsize=8)
    a.set_xticks(xs)
    a.set_xticklabels(["%s\n%.0f to %.0f" % (lab, v['lo'], v['hi']) for lab, v in zip(("Q1", "Q2", "Q3", "Q4"), q)])
    a.set_xlabel("Week-4 pre-dose concentration quartile (ug/mL), 600 mg")
    a.set_ylabel("Week-12 clinical remission (%)")
    a.set_ylim(0, 70)
    a.axhline(F['fc'][300], color=TEAL, lw=1.6, zorder=3)
    a.text(-0.42, 64, "Forecast remission (charter 4):  300 mg %.1f%%   600 mg %.1f%%   900 mg %.1f%%"
           % (F['fc'][300], F['fc'][600], F['fc'][900]), fontsize=8.6, color=TEAL, va="center")
    a.text(1.5, F["fc"][300] + 1.4, "300 mg forecast", ha="center", fontsize=8.2, color=TEAL)
    a.set_title("Remission is flat across the exposure the dose sets", loc="left", fontsize=10.5, color=INK, fontweight="bold")
    a.spines[["top", "right"]].set_visible(False)
    a.grid(axis="y", color="#e5e7eb", lw=0.8, zorder=0)
    vis = ["Week 4\n(pre-dose)", "Week 8\n(pre-dose)", "Week 12"]
    ratios = [F['gmr']['tr4'], F['gmr']['tr8'], F['gmr']['c12']]
    b.plot(range(3), ratios, color=SAND, lw=2, marker="o", ms=8, zorder=3)
    for x, v in enumerate(ratios):
        b.text(x, v + 0.18, "%.2f" % v, ha="center", color=INK, fontsize=9)
    b.axhline(1.0, color=GREY, lw=1, ls="--")
    b.text(2.25, 1.06, "no difference", ha="right", fontsize=8, color=GREY)
    b.set_xticks(range(3))
    b.set_xticklabels(vis)
    b.set_ylim(0, 4)
    b.set_xlim(-0.4, 2.4)
    b.set_ylabel("Remitters to non-remitters, geometric mean ratio")
    b.set_title("Concentrations part only after healing starts", loc="left", fontsize=10.5, color=INK, fontweight="bold")
    b.spines[["top", "right"]].set_visible(False)
    b.grid(axis="y", color="#e5e7eb", lw=0.8, zorder=0)
    fig.text(0.01, 0.01, "Source: CVN2184-UC-201 database lock 14 July 2026 and Brightwater bioanalytical results; "
             "%d of %d dovikizumab patients with a valid week-4 pre-dose sample." % (F['n'], F['n_arm']),
             fontsize=7.6, color=GREY)
    fig.tight_layout(rect=(0, 0.04, 1, 1))
    fig.savefig(path, dpi=150, metadata={"Software": None})
    plt.close(fig)


THIN = Side(style="thin", color="B7BEC8")


def _sheet(ws, header, rows, widths, fmt=None):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="14505C")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(r)
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "B2"
    if fmt:
        for col, f in fmt.items():
            for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                for c in row:
                    if isinstance(c.value, (int, float)):
                        c.number_format = f
                        c.alignment = Alignment(horizontal="right")


def tables(path, F):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dose forecasts"
    _sheet(ws, ["Induction regimen", "Forecast week-12 clinical remission (%)", "Points below the highest forecast",
                "Within 5 points (charter 5)"],
           [["%d mg IV weeks 0, 4, 8" % d, round(F['fc'][d], 1), round(max(F['fc'].values()) - F['fc'][d], 1),
             "yes" if max(F['fc'].values()) - F['fc'][d] <= K.THR else "no"] for d in K.DOSES],
           [26, 22, 22, 18], {2: "0.0", 3: "0.0"})
    ws.append([])
    ws.append(["Proposed phase 3 induction dose: %d mg (the lowest candidate within 5 points of the highest forecast)." % F['call']])
    ws.append(["Exposure-response analysis set: %d of %d dovikizumab patients with a valid week-4 pre-dose concentration; "
               "observed week-12 remission in the set %.1f%%." % (F['n'], F['n_arm'], F['obs_set'])])
    S = wb.create_sheet("Safety")
    A = F['ask_a']
    _sheet(S, ["Treatment-emergent (SAP 6.3 to 6.6)", "Dovikizumab 600 mg", "Placebo"],
           [["Patients randomised", A['dovikizumab']['n'], A['placebo']['n']],
            ["Patients with at least one serious infection", A['dovikizumab']['pts_si'], A['placebo']['pts_si']],
            ["Patients with at least one infusion-related reaction", A['dovikizumab']['pts_irr'], A['placebo']['pts_irr']],
            ["Serious infection events", A['dovikizumab']['events_si'], A['placebo']['events_si']],
            ["Exposure (patient-years)", round(A['dovikizumab']['py'], 1), round(A['placebo']['py'], 1)],
            ["Serious infections per 100 patient-years", round(A['dovikizumab']['rate_si'], 1), round(A['placebo']['rate_si'], 1)]],
           [52, 20, 14], {2: "#,##0.0", 3: "#,##0.0"})
    for row in S.iter_rows(min_row=2, max_row=5, min_col=2, max_col=3):
        for c in row:
            c.number_format = "#,##0"
    I = wb.create_sheet("Immunogenicity")
    B = F['ask_b']
    rows = []
    for vis in ("Week 4", "Week 8", "Week 12"):
        rows.append([vis, B['dovikizumab'][vis]['te'], B['dovikizumab'][vis]['n'], round(B['dovikizumab'][vis]['pct'], 1),
                     B['placebo'][vis]['te'], B['placebo'][vis]['n'], round(B['placebo'][vis]['pct'], 1)])
    _sheet(I, ["Visit", "Dovikizumab: TE-ADA positive", "Dovikizumab: evaluable", "Dovikizumab: incidence (%)",
               "Placebo: TE-ADA positive", "Placebo: evaluable", "Placebo: incidence (%)"], rows,
           [12, 16, 14, 16, 14, 12, 14], {4: "0.0", 7: "0.0"})
    C = wb.create_sheet("Symptomatic remission")
    Cc = F['ask_c']
    _sheet(C, ["Reference day (SAP 6.10)", "Dovikizumab: in symptomatic remission", "Dovikizumab: in study",
               "Placebo: in symptomatic remission", "Placebo: in study"],
           [["Week %d (day %d)" % (w, dict(K.DIARY_WEEKS)[w] + 1), Cc['dovikizumab'][w]['rem'], Cc['dovikizumab'][w]['n'],
             Cc['placebo'][w]['rem'], Cc['placebo'][w]['n']] for w, _ in K.DIARY_WEEKS], [22, 20, 16, 18, 14])
    N = wb.create_sheet("Notes")
    for line in ("Forecasts: logistic regression of week-12 clinical remission (SAP 6.9, non-responder imputation) on the log "
                 "week-4 pre-dose concentration with the randomisation stratum, exposure at 300 and 900 mg by proportional "
                 "scaling (investigator's brochure 5.2), averaged over the exposure-response analysis set (charter 4).",
                 "Week-4 pre-dose samples drawn after the infusion started (%d in the dovikizumab arm, laboratory UTC against the eCRF's local start) are "
                 "not pre-dose samples (SAP 6.8) and were not used." % F['late'],
                 "Adverse events: the eCRF log re-enters an ongoing event at each visit's review; each event is the "
                 "patient's chain of records of one preferred term from review to review, with the onset entered at its "
                 "latest review (sites corrected onsets first entered as the visit or admission date) and serious if any "
                 "record is; partial onset dates follow SAP 6.2.",
                 "ADA: each draw is split into two aliquot tubes; a draw is the patient's tubes collected on one date, "
                 "whatever visit the shipment manifest gives, and each tier's result is its latest in a run that met "
                 "acceptance in Brightwater's run log.",
                 "E-diary: entries are the patient's who held the handset on the diary date. Patients of the four sites "
                 "that closed kept their handsets and were registered with the e-diary vendor at the receiving site under "
                 "a new subject number; their entries are counted under their randomised number. A patient's value for a "
                 "day is the patient's last entry for it.",
                 "Prepared by Clinical Pharmacology for the end-of-phase-2 briefing package, October 2026."):
        N.append([line])
    N.column_dimensions["A"].width = 130
    for row in N.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def memo(path, F):
    import docx
    from docx.shared import Pt, RGBColor
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    h = d.add_paragraph()
    r = h.add_run("Corvane Therapeutics | Clinical Pharmacology")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x5a, 0x62, 0x70)
    t = d.add_paragraph()
    r = t.add_run("Phase 3 induction dose for dovikizumab in ulcerative colitis: %d mg IV at weeks 0, 4 and 8" % F['call'])
    r.bold = True
    r.font.size = Pt(14)
    d.add_paragraph("To: Clinical Development Committee   From: Head of Clinical Pharmacology   Date: October 2026   "
                    "For: end-of-phase-2 briefing package, section 4.2 (due 4 December 2026)")
    d.add_paragraph(
        "We propose %d mg. Under the dose-selection charter (CVN-CP-0142) the forecast week-12 clinical remission rate is "
        "%.1f per cent at 300 mg, %.1f per cent at 600 mg and %.1f per cent at 900 mg; 300 mg sits %.1f points under the "
        "highest forecast, inside the charter's 5 points, and is the lowest candidate there."
        % (F['call'], F['fc'][300], F['fc'][600], F['fc'][900], max(F['fc'].values()) - F['fc'][300]))
    d.add_heading("Why the forecast rests on the week-4 pre-dose concentration", level=2)
    d.add_paragraph(
        "Phase 2b tested one dose, so the exposure-response comes from the spread of exposure within the 600 mg arm. "
        "Which concentration carries the dose's effect is the question that decides the forecast. At week 4, before "
        "most patients' colitis had begun to heal, remitters and non-remitters had the same concentrations (geometric "
        "mean ratio %.2f). By week 8 remitters ran %.2f times higher and by week 12 %.2f times higher. Antibody clearance "
        "falls as the inflamed bowel heals, so the later concentrations are partly a result of remission rather than its "
        "cause: an exposure-response fitted to them is steep (300 mg about %.0f per cent, 900 mg about %.0f) and would "
        "carry 900 mg, but the steepness is the healing, not the dose." % (F['gmr']['tr4'], F['gmr']['tr8'], F['gmr']['c12'],
                                                                          F['late_300'], F['late_900']))
    d.add_paragraph(
        "On the week-4 concentration, which the dose sets before healing can move it, week-12 remission is flat across "
        "the quartiles of exposure (%s per cent), and every candidate dose sits on the plateau of the dose's effect. The "
        "forecast holds with or without the randomisation stratum and baseline weight and modified Mayo score in the "
        "model (300 mg %.1f to %.1f per cent)." % (", ".join("%.0f" % v['rate'] for v in F['quartiles']), F['std_lo'], F['std_hi']))
    d.add_heading("Records and conventions", level=2)
    d.add_paragraph(
        "%d of the %d dovikizumab patients have a valid week-4 pre-dose concentration; %d week-4 samples drawn after the infusion "
        "had started were set aside under SAP 6.8 by comparing the laboratory's UTC collection time with the site's local "
        "infusion start. Remission is the database-lock flag (central adjudicated endoscopy, non-responder imputation). "
        "Safety, immunogenicity and e-diary summaries are in phase2b_induction_tables.xlsx." % (F['n'], F['n_arm'], F['late']))
    tb = d.add_table(rows=1, cols=4)
    tb.style = "Light Grid Accent 1"
    for c, v in zip(tb.rows[0].cells, ("Regimen", "Forecast remission", "Below highest", "Within 5 points")):
        c.text = v
    for dose in K.DOSES:
        row = tb.add_row().cells
        row[0].text = "%d mg" % dose
        row[1].text = "%.1f%%" % F['fc'][dose]
        row[2].text = "%.1f" % (max(F['fc'].values()) - F['fc'][dose])
        row[3].text = "yes" if max(F['fc'].values()) - F['fc'][dose] <= K.THR else "no"
    d.add_paragraph("Figure: exposure_response_chart.png. Forecasts compared unrounded (charter 5).").runs[0].italic = True
    d.save(path)
