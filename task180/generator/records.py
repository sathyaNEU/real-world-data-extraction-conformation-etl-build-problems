"""Writers for the pack's data files, each in its source system's own layout: the survey's catalogues and station
records, the plants' telemetry and the operators' monthly reports as the data office keeps them."""
import csv
import datetime as dt
import json

import numpy as np
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq

import sim as S
import world as W

CAT = "seismic_catalogue_reviewed_2023-01_2026-09.csv"
AUTO = "origins_2023-01_2026-09.csv"
HOURLY = "injection_hourly_2023-01_2026-09.parquet"
REPORTS = "injection_reports_2023-2026.xlsx"
INVENTORY = "station_inventory_2026-10-12.xml"
GAPS = "availability_gaps_2026-01_2026-09.csv"
VISITS = "field_visits_2026.xlsx"
PLANT_REG = "plant_register.json"
WELL_REG = "well_register.csv"
FEDERAL = "federal_catalogue_2023-2026.csv"
GW = "groundwater_levels_2025-01_2026-09.csv"

UTC = dt.timezone.utc


def iso_utc(t, digits=2):
    """Epoch seconds as an ISO time in UTC with the given decimals of a second."""
    t = float(t)
    whole = int(np.floor(t))
    frac = t - whole
    s = dt.datetime.fromtimestamp(whole, UTC).strftime("%Y-%m-%dT%H:%M:%S")
    if digits:
        f = int(round(frac * 10 ** digits))
        if f == 10 ** digits:
            s = dt.datetime.fromtimestamp(whole + 1, UTC).strftime("%Y-%m-%dT%H:%M:%S")
            f = 0
        s += ".%0*d" % (digits, f)
    return s + "Z"


# ------------------------------------------------------------------------------------------------- catalogues
CAT_COLS = ["event_id", "origin_time_utc", "latitude", "longitude", "depth_km", "ml", "ml_n_sta", "n_phases",
            "az_gap_deg", "rms_s", "h_err_km", "depth_fixed", "event_type", "region", "analyst", "reviewed_utc"]


def write_catalogue(path, cat):
    c = cat.sort_values("t").reset_index(drop=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(CAT_COLS)
        for e in c.itertuples():
            w.writerow([e.event_id, iso_utc(e.t, 2), "%.4f" % e.lat, "%.4f" % e.lon, "%.2f" % e.depth_km, "%.2f" % e.ml,
                        int(e.ml_n_sta), int(e.n_phases), int(e.az_gap_deg), "%.3f" % e.rms_s, "%.2f" % e.h_err_km,
                        e.depth_fixed, e.event_type, e.region, e.analyst, iso_utc(e.reviewed_at, 0)])
    return len(c)


AUTO_COLS = ["origin_id", "creation_time_utc", "origin_time_utc", "latitude", "longitude", "depth_km", "ml", "n_picks",
             "az_gap_deg", "author"]


def write_automatic(path, auto):
    a = auto.sort_values("created").reset_index(drop=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(AUTO_COLS)
        for o in a.itertuples():
            w.writerow([o.origin_id, iso_utc(o.created, 3), iso_utc(o.t, 2), "%.4f" % o.lat, "%.4f" % o.lon,
                        "%.1f" % o.depth_km, "%.2f" % o.ml_auto, int(o.n_picks), int(round(o.az_gap)), o.author])
    return len(a)


def write_federal(path, fed):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["evid", "time_utc", "lat", "lon", "depth_km", "mag", "mag_type"])
        for e in fed.itertuples():
            w.writerow([e.evid, iso_utc(e.t, 1), "%.2f" % e.lat, "%.2f" % e.lon, "%.1f" % e.depth_km, "%.1f" % e.mag,
                        e.mag_type])
    return len(fed)


# ------------------------------------------------------------------------------------------------- injection
def write_hourly(path, hourly):
    h = hourly.sort_values(["well", "t"]).reset_index(drop=True)
    tbl = pa.table({
        "plant": pa.array(h["plant"].to_numpy(), pa.string()),
        "well": pa.array(h["well"].to_numpy(), pa.string()),
        "time_utc": pa.array(h["t"].to_numpy().astype("int64"), pa.timestamp("s", tz="UTC")),
        "flow_l_s": pa.array(np.round(h["flow_ls"].to_numpy(), 2), pa.float64()),
        "whp_bar": pa.array(np.round(h["pressure_bar"].to_numpy(), 2), pa.float64()),
        "temperature_c": pa.array(np.round(h["temp_c"].to_numpy(), 1), pa.float64()),
    }).replace_schema_metadata({})
    pq.write_table(tbl, path, compression="zstd", row_group_size=65536, use_dictionary=["plant", "well"],
                   write_statistics=True)
    return tbl.num_rows


LEGACY_ID = {"MST-GT4": "MST-Th4a"}     # the operator still reports its second injector under the drilling name


def report_rows(inj):
    """The operators' monthly reports, one row per well and month as each was received. To December 2025 they came
    by e-mail and the data office filed each under the month it reports; from January 2026 they come through the
    authority's reporting portal, whose export files each report under the month it was submitted in."""
    r = S.rng(16)
    rows = []
    for x in inj.sort_values(["year", "month", "plant", "role", "well"]).itertuples():
        y, m = int(x.year), int(x.month)
        ny, nm = (y + 1, 1) if m == 12 else (y, m + 1)
        rec = dt.date(ny, nm, int(r.integers(2, 11)))
        base = dict(operator=W.OPERATOR[x.plant], well=LEGACY_ID.get(x.well, x.well),
                    function="injection" if x.role == "injection" else "production", volume_m3=int(x.volume_reported),
                    whp_bar=round(float(x.pressure_bar), 1) if x.role == "injection" else None,
                    temp_c=round(float(x.temp_c), 1), received=rec, op_year=y, op_month=m)
        if (y, m) < (2025, 12):
            rows.append(dict(base, sheet=str(y), month="%d-%02d" % (y, m)))
        elif (y, m) == (2025, 12):
            # sent both ways in the changeover month
            rows.append(dict(base, sheet="2025", month="2025-12"))
            rows.append(dict(base, sheet="2026", month="%d-%02d" % (ny, nm),
                             received=dt.date(ny, nm, min(10, rec.day + int(r.integers(0, 3))))))
        else:
            rows.append(dict(base, sheet="2026", month="%d-%02d" % (ny, nm)))
    return rows


def write_reports(path, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    wb.remove(wb.active)
    head = ["month", "operator", "well", "function", "volume_m3", "whp_bar", "temperature_c", "received"]
    widths = [9, 34, 11, 11, 12, 9, 8, 11]
    for sheet in ("2023", "2024", "2025", "2026"):
        ws = wb.create_sheet(sheet)
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
        sel = sorted((x for x in rows if x["sheet"] == sheet), key=lambda x: (x["month"], x["operator"], x["well"]))
        for x in sel:
            ws.append([x["month"], x["operator"], x["well"], x["function"], x["volume_m3"], x["whp_bar"], x["temp_c"],
                       x["received"]])
            ws.cell(ws.max_row, 5).number_format = "#,##0"
            ws.cell(ws.max_row, 6).number_format = "0.0"
            ws.cell(ws.max_row, 7).number_format = "0.0"
            ws.cell(ws.max_row, 8).number_format = "yyyy-mm-dd"
            ws.cell(ws.max_row, 8).alignment = Alignment(horizontal="right")
        for i, wd in enumerate(widths):
            ws.column_dimensions["ABCDEFGH"[i]].width = wd
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = "A1:H%d" % ws.max_row
    wb.save(path)
    return sum(1 for x in rows)


# ------------------------------------------------------------------------------------------------- registers
def write_plant_register(path, st):
    plants = []
    for p in W.PLANTS:
        net = sorted({e["station"] for e in st if e["plant"] == p and (e["end"] is None or e["loc"] == "10"
                                                                         or e["end"] > W.RECORD_END)})
        plants.append({
            "plant": p, "name": W.NAME[p], "operator": W.OPERATOR[p], "municipality": W.NAME[p],
            "wellhead": {"lat": W.WELLHEAD[p][0], "lon": W.WELLHEAD[p][1]},
            "injection_wells": [w for w, _ in S.INJ_WELLS[p]], "production_wells": S.PROD_WELLS[p],
            "reinjection_temperature_c": W.REINJECTION_TEMP_C[p], "monitoring_network": net,
        })
    reg = {"register": "Tiefe Geothermie: Anlagen unter seismischer Überwachung", "kept_by": W.SURVEY + ", Seismologischer Dienst",
           "as_of": "2026-10-01", "network_code": W.NET, "plants": plants}
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(reg, fh, indent=2, ensure_ascii=False)
        fh.write("\n")


PROD_HOLE = {
    "MST-GT1": ((47.9262, 11.4101, 2950), (47.9318, 11.4204, 3540)), "MST-GT3": ((47.9175, 11.3968, 3010), (47.9126, 11.3861, 3600)),
    "PFB-GT1": ((48.0841, 11.7088, 2870), (48.0877, 11.6956, 3440)), "LZR-GT1": ((47.9871, 11.5990, 3120), (47.9942, 11.5917, 3690)),
    "WDR-GT1": ((47.8541, 11.6154, 3290), (47.8478, 11.6082, 3860)), "OTM-GT1": ((48.0151, 11.9128, 2960), (48.0093, 11.9061, 3510)),
}
FORMER = {"MST-GT4": "MST-Th4a", "MST-GT3": "MST-B3"}     # designations before the March 2024 register revision
SPUD = {"MST-GT1": "2014-03-11", "MST-GT2": "2014-07-02", "MST-GT3": "2018-05-23", "MST-GT4": "2018-10-15",
        "PFB-GT1": "2016-09-05", "PFB-GT2": "2017-01-19", "LZR-GT1": "2015-04-27", "LZR-GT2": "2015-09-08",
        "WDR-GT1": "2017-06-12", "WDR-GT2": "2017-10-30", "OTM-GT1": "2018-02-06", "OTM-GT2": "2018-06-18"}


def write_well_register(path):
    cols = ["well_id", "former_designation", "plant", "function", "spud_date", "wellhead_lat", "wellhead_lon", "oh_top_lat",
            "oh_top_lon", "oh_top_tvd_m", "oh_bottom_lat", "oh_bottom_lon", "oh_bottom_tvd_m"]
    rows = []
    for p in W.PLANTS:
        for k, (w, _) in enumerate(S.INJ_WELLS[p]):
            rows.append((w, p, "injection", W.OPEN_HOLE[p][k]))
        for w in S.PROD_WELLS[p]:
            rows.append((w, p, "production", PROD_HOLE[w]))
    rows.sort(key=lambda x: x[0])
    with open(path, "w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh, lineterminator="\n")
        wr.writerow(cols)
        for w, p, f, ((la1, lo1, d1), (la2, lo2, d2)) in rows:
            wr.writerow([w, FORMER.get(w, ""), p, f, SPUD[w], "%.4f" % W.WELLHEAD[p][0], "%.4f" % W.WELLHEAD[p][1],
                         "%.4f" % la1, "%.4f" % lo1, d1, "%.4f" % la2, "%.4f" % lo2, d2])
    return len(rows)


# ------------------------------------------------------------------------------------------------- stations
SENSOR = "Lennartz LE-3D/1s"
OLD_LOGGER, NEW_LOGGER = "Earth Data PR6-24", "Nanometrics Centaur CTR4-6S"


def _x(d):
    return d.astimezone(UTC).strftime("%Y-%m-%dT%H:%M:%SZ") if d is not None else None


def write_inventory(path, st):
    by = {}
    for e in st:
        by.setdefault(e["station"], []).append(e)
    lines = ['<?xml version="1.0" encoding="UTF-8"?>',
             '<FDSNStationXML xmlns="http://www.fdsn.org/xml/station/1" schemaVersion="1.2">',
             "  <Source>LDG</Source>", "  <Sender>%s, Seismologischer Dienst</Sender>" % W.SURVEY,
             "  <Created>2026-10-12T05:48:17Z</Created>",
             '  <Network code="%s" startDate="2016-05-18T00:00:00Z">' % W.NET,
             "    <Description>Induced-seismicity networks of the deep geothermal plants, southern Molasse basin</Description>"]
    for code, plant, off, site in S.STATIONS:
        eps = sorted(by[code], key=lambda e: e["start"])
        s0, s1 = eps[0]["start"], eps[-1]["end"]
        e0 = eps[0]
        end_attr = (' endDate="%s"' % _x(s1)) if s1 is not None else ""
        lines.append('    <Station code="%s" startDate="%s"%s>' % (code, _x(s0), end_attr))
        lines += ["      <Latitude>%.5f</Latitude>" % e0["lat"], "      <Longitude>%.5f</Longitude>" % e0["lon"],
                  "      <Elevation>%.1f</Elevation>" % e0["elev"], "      <Site>", "        <Name>%s</Name>" % site,
                  "      </Site>"]
        for e in eps:
            logger = NEW_LOGGER if e["loc"] == "10" else OLD_LOGGER
            for cha, az, dip in (("HHZ", 0, -90), ("HHN", 0, 0), ("HHE", 90, 0)):
                ea = (' endDate="%s"' % _x(e["end"])) if e["end"] is not None else ""
                lines += ['      <Channel code="%s" locationCode="%s" startDate="%s"%s>' % (cha, e["loc"], _x(e["start"]), ea),
                          "        <Latitude>%.5f</Latitude>" % e["lat"], "        <Longitude>%.5f</Longitude>" % e["lon"],
                          "        <Elevation>%.1f</Elevation>" % e["elev"], "        <Depth>0.0</Depth>",
                          "        <Azimuth>%d</Azimuth>" % az, "        <Dip>%d</Dip>" % dip,
                          "        <SampleRate>200</SampleRate>",
                          "        <Sensor>", "          <Type>%s</Type>" % SENSOR, "        </Sensor>",
                          "        <DataLogger>", "          <Type>%s</Type>" % logger, "        </DataLogger>",
                          "      </Channel>"]
        lines.append("    </Station>")
    lines += ["  </Network>", "</FDSNStationXML>", ""]
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))
    return len(S.STATIONS)


def write_gaps(path, g):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["net", "sta", "loc", "cha", "gap_start", "gap_end", "gap_s"])
        for x in g.itertuples():
            w.writerow([W.NET, x.station, x.loc, x.channel, iso_utc(x.s, 0), iso_utc(x.e, 0), int(x.e - x.s)])
    return len(g)


WORK = {"exchange": "datalogger swap", "service": None}
SERVICE_WORK = ["battery check, enclosure resealed", "GPS antenna cable replaced", "battery replaced",
                "vault drained, sensor levelled", "modem replaced", "solar controller replaced"]


def write_visits(path, field):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    r = S.rng(17)
    wb = Workbook()
    ws = wb.active
    ws.title = "Visits 2026"
    ws.append(["date", "station", "site", "arrival", "departure", "work", "by"])
    for c in ws[1]:
        c.font = Font(bold=True)
    sites = {code: site for code, _, _, site in S.STATIONS}
    for v in sorted(field, key=lambda v: v["arrival"]):
        if v["arrival"].date() > W.EXTRACT_DATE:
            continue
        work = WORK["exchange"] if v["kind"] == "exchange" else SERVICE_WORK[int(r.integers(0, len(SERVICE_WORK)))]
        ws.append([v["arrival"].date(), v["station"], sites[v["station"]], v["arrival"].strftime("%H:%M"),
                   v["departure"].strftime("%H:%M"), work, "E. Scheel"])
        ws.cell(ws.max_row, 1).number_format = "dd.mm.yyyy"
    for col, wd in zip("ABCDEFG", (11, 8, 34, 8, 10, 58, 10)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A2"
    wb.save(path)
    return ws.max_row - 1


def write_groundwater(path, gw):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["plant", "well", "month", "level_m_bgl", "temperature_c", "ec_us_cm"])
        for x in gw.itertuples():
            w.writerow([x.plant, x.well, "%d-%02d" % (x.year, x.month), "%.2f" % x.level_m_bgl, "%.1f" % x.temp_c, x.ec_us_cm])
    return len(gw)
