"""The goldens, computed from target/ only: the 2027 outlook memo to the mining authority (amber_outlook_2027.docx, with
its chart) and the 2026 annex workbook (tls_annex_2026.xlsx), and generator/figures.json with every figure the
write-up states."""
import io
import json
import os
import sys

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt   # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import asks as A           # noqa: E402
import models as M         # noqa: E402
import postprocess as PP   # noqa: E402
import world as W          # noqa: E402

DOCX = "amber_outlook_2027.docx"
XLSX = "tls_annex_2026.xlsx"
LEAD = "Gundula Atzler"
TOL = "8"                       # every rate convention on the whole record keeps each 2027 cell within this (checks.py)
INK, INK2, GRID, PAPER = "#1d2329", "#56606a", "#dfe3e7", "#ffffff"
PFB_C, OTHER_C = "#b5401f", ["#7d8a96", "#9aa6b1", "#b6c0c9", "#d0d7dd"]
QLAB = {1: "Q1", 2: "Q2", 3: "Q3", 4: "Q4"}


def compute(target):
    d = M.load(target)
    c = M.counted(d)
    r = M.rates(c, d)
    r_all = M.rates(c, d, quiet_plants=())
    g = M.forecast_from_rates(r, M.plan_reports(d))
    pfb = c[c["plant"] == "PFB"]
    t0 = M.table(c)
    return dict(d=d, c=c, g=g, rate=r, rate_all=r_all, n_all=int(len(pfb)), n_quiet=int(M.quiet(pfb["h"].to_numpy()).sum()),
                located=t0, asks=A.golden(target), over=M.over(g))


def whole(x):
    return int(round(x))


def figures(c):
    g = c["g"]
    others = max(((p, q) for p in W.PLANTS if p != "PFB" for q in (1, 2, 3, 4)), key=lambda k: g[k[0]][k[1]])
    a = c["asks"]
    out = {
        "plants_over_amber": [W.NAME[p] for p in sorted({p for p, _ in c["over"]})],
        "quarters_over_amber": ["%s %s 2027" % (W.NAME[p], QLAB[q]) for p, q in c["over"]],
        "amber_more_than": W.AMBER,
        "forecast_2027": {W.NAME[p]: {QLAB[q]: whole(g[p][q]) for q in (1, 2, 3, 4)} for p in W.PLANTS},
        "highest_other": {"plant": W.NAME[others[0]], "quarter": QLAB[others[1]], "count": whole(g[others[0]][others[1]])},
        "count_ml": "1.0",
        "per_volume_m3": "100,000",
        "quiet_hours_per_day": "5.5",
        "highest_located_2026Q1": {"plant": W.NAME[max(W.PLANTS, key=lambda p: c["located"][p][(2026, 1)])],
                                   "count": max(c["located"][p][(2026, 1)] for p in W.PLANTS)},
        "pfb_located_2026Q1": c["located"]["PFB"][(2026, 1)],
        "pfb_located_2025Q4": c["located"]["PFB"][(2025, 4)],
        "pfb_located_record": c["n_all"],
        "pfb_located_quiet_hours": c["n_quiet"],
        "pfb_rate_per_100k_m3_quiet": "%.1f" % (c["rate"]["PFB"] * 1e5),
        "pfb_rate_per_100k_m3_all_hours": "%.1f" % (c["rate_all"]["PFB"] * 1e5),
        "pfb_rate_ratio": "%.1f" % (c["rate"]["PFB"] / c["rate_all"]["PFB"]),
        "quiet_hours_local": "23:30 to 05:00",
        "cell_tolerance_pct": TOL,
        "design_completeness_ml": "0.6",
        "annex_2026": {W.NAME[p]: {"Q%d" % q: {"volume_m3": "{:,}".format(a[p][(y, q)]["volume_m3"]),
                                               "automatic": a[p][(y, q)]["automatic"],
                                               "hours_under_four": "%.1f" % a[p][(y, q)]["hours_under_four"]}
                                   for y, q in A.QS} for p in W.PLANTS},
    }
    return out


# ------------------------------------------------------------------------------------------------- the chart
def chart_png(c):
    g = c["g"]
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9, "svg.hashsalt": "t180",
                         "axes.edgecolor": INK2, "axes.labelcolor": INK2, "xtick.color": INK, "ytick.color": INK2})
    fig, ax = plt.subplots(figsize=(7.2, 3.9), dpi=170)
    fig.patch.set_facecolor(PAPER)
    ax.set_facecolor(PAPER)
    width = 0.155
    order = ["PFB", "MOS", "LZR", "OTM", "WDR"]
    for i, p in enumerate(order):
        xs = [q - 1 + (i - 2) * width for q in (1, 2, 3, 4)]
        ys = [g[p][q] for q in (1, 2, 3, 4)]
        col = PFB_C if p == "PFB" else OTHER_C[i - 1]
        ax.bar(xs, ys, width=width * 0.92, color=col, label=W.NAME[p], zorder=3)
        if p == "PFB":
            for x, y in zip(xs, ys):
                ax.text(x, y + 3, "%d" % whole(y), ha="center", va="bottom", fontsize=8, color=PFB_C, fontweight="bold")
    ax.axhline(W.AMBER, color="#d08a00", linewidth=1.4, zorder=4)
    ax.text(3.55, W.AMBER + 3, "amber: more than %d" % W.AMBER, ha="right", va="bottom", fontsize=8, color="#9a6500")
    ax.set_xticks([0, 1, 2, 3])
    ax.set_xticklabels(["Q1 2027", "Q2 2027", "Q3 2027", "Q4 2027"])
    ax.set_ylabel("Earthquakes of ML 1.0 or more within 5 km")
    ax.yaxis.grid(True, color=GRID, linewidth=0.8, zorder=0)
    ax.set_ylim(0, 250)
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.tick_params(axis="y", length=0)
    ax.tick_params(axis="x", length=0)
    ax.legend(ncol=5, frameon=False, fontsize=8, loc="upper center", bbox_to_anchor=(0.5, -0.09))
    fig.suptitle("Only Pfeilbrunn goes over amber, in the first and fourth quarters", x=0.015, ha="left", fontsize=10.5,
                 fontweight="bold", color=INK)
    fig.tight_layout(rect=(0, 0, 1, 0.95))
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor=PAPER, metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


# ------------------------------------------------------------------------------------------------- the memo
def write_docx(path, c, fig):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.3)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.font.size = Pt(10)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.12

    def para(text="", bold=False, size=None, color=None, after=None, italic=False, align=None):
        p = doc.add_paragraph()
        if text:
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        if align:
            p.alignment = align
        return p

    def no_borders(tbl):
        tblPr = tbl._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement("w:%s" % e)
            el.set(qn("w:val"), "nil")
            b.append(el)
        tblPr.append(b)

    para("Landesdienst für Geologie  |  Seismologischer Dienst", size=8.5, color="56606A", after=10)
    head = doc.add_table(rows=5, cols=2)
    no_borders(head)
    head.autofit = False
    head.columns[0].width, head.columns[1].width = Cm(2.6), Cm(13.8)
    for i, (k, v) in enumerate([("To", "%s, Referat Geothermie" % W.AUTHORITY),
                                ("From", "Seismologischer Dienst, induced seismicity (%s)" % LEAD),
                                ("Date", "12 October 2026"), ("Reference", "LDG-SD-AO-2026-07"),
                                ("Subject", "Traffic-light outlook for 2027 under the borehole networks")]):
        head.cell(i, 0).width, head.cell(i, 1).width = Cm(2.6), Cm(13.8)
        r = head.cell(i, 0).paragraphs[0].add_run(k)
        r.bold = True
        r.font.size = Pt(9)
        r2 = head.cell(i, 1).paragraphs[0].add_run(v)
        r2.font.size = Pt(9)
        for cell in head.rows[i].cells:
            cell.paragraphs[0].paragraph_format.space_after = Pt(1)
    para("", after=4)

    g = c["g"]
    f = figures(c)
    q1, q4 = whole(g["PFB"][1]), whole(g["PFB"][4])
    q2, q3 = whole(g["PFB"][2]), whole(g["PFB"][3])
    ho = f["highest_other"]
    para("Pfeilbrunn will run over amber in the first and fourth quarters of 2027", bold=True, size=13, after=8)
    para("On the networks that make the counts from 1 January 2027 we expect Pfeilbrunn to locate about %d earthquakes "
         "of ML 1.0 or more within 5 km of its injection well in the first quarter and about %d in the fourth, against "
         "the amber level of more than %d. Its second and third quarters stay under the line (about %d and %d), so amber "
         "falls in the heating season. No other plant comes near amber in any quarter; the highest is %s at about %d "
         "in the first quarter." % (q1, q4, W.AMBER, q2, q3, ho["plant"], ho["count"]))
    para("Pfeilbrunn has never been close to amber on its present network: %d counted in the first quarter of 2026 and "
         "%d in the fourth quarter of 2025. Those counts are what its town stations could record, not what the "
         "reservoir produces. The four stations stand in the town (school basement, waterworks, fire station, building "
         "yard), and through the working day traffic and works hide all but the larger earthquakes. From 23:30 to "
         "05:00 they record every earthquake down to about ML 0.6, as the rural networks do at every hour. Over January "
         "2023 to September 2026 they located %d of the plant's earthquakes of ML 1.0 or more in those five and a half "
         "hours of the night, against %d over the whole day. At the night rate round the clock the reservoir produces "
         "%s such earthquakes per 100,000 m³ injected, %s times the %s its full-day record shows. The borehole "
         "stations, designed to locate every earthquake of ML %s or more within 5 km, will count at that rate."
         % (f["pfb_located_2026Q1"], f["pfb_located_2025Q4"], c["n_quiet"], c["n_all"], f["pfb_rate_per_100k_m3_quiet"],
            f["pfb_rate_ratio"], f["pfb_rate_per_100k_m3_all_hours"], f["design_completeness_ml"]))
    para("The usual completeness check does not show this. Pfeilbrunn's catalogue peaks near ML 0.4 like the others "
         "because the night fills the small magnitudes; what is missing is the day, and only the record by hour shows it.")
    para("Under section 4.2 of the scheme each amber notice obliges Stadtwerke Pfeilbrunn to cut injection by at least a "
         "quarter within 48 hours; its 2027 plan, filed on 1 September, sets every quarter at the 2026 volume and does not "
         "yet provide for that.")
    para("The rural networks at Moorstetten, Wendelsried and Ottenmoos record every earthquake above about ML 0.6 at any "
         "hour, and Lanzenried has done so since LZR5 replaced LZR4 nearer the well on 15 July 2025. Their 2027 counts "
         "follow their own rates; Lanzenried's higher counts since mid-2025 are that station move, not a change in the "
         "reservoir.")
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(io.BytesIO(fig), width=Cm(16.2))
    pic.paragraph_format.space_after = Pt(2)
    pic.paragraph_format.keep_with_next = True
    para("Figure 1. Expected earthquakes of ML 1.0 or more within 5 km of the injection wells, each quarter of 2027, "
         "against the amber level.", size=8, color="56606A", after=10)

    hp = para("Expected count per quarter, 2027", bold=True, size=10, after=3)
    hp.paragraph_format.keep_with_next = True
    tb = doc.add_table(rows=6, cols=5)
    tb.style = "Table Grid"
    tb.alignment = WD_TABLE_ALIGNMENT.LEFT
    tb.autofit = False
    tb.columns[0].width = Cm(4.2)
    for k in range(1, 5):
        tb.columns[k].width = Cm(2.2)
    hdr = ["Plant", "Q1", "Q2", "Q3", "Q4"]
    for j, h in enumerate(hdr):
        r = tb.cell(0, j).paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9)
        if j:
            tb.cell(0, j).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for i, p in enumerate(["PFB", "MOS", "LZR", "OTM", "WDR"], 1):
        r = tb.cell(i, 0).paragraphs[0].add_run(W.NAME[p])
        r.font.size = Pt(9)
        for q in (1, 2, 3, 4):
            cell = tb.cell(i, q).paragraphs[0]
            cell.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            rr = cell.add_run("%d" % whole(g[p][q]))
            rr.font.size = Pt(9)
            if g[p][q] > W.AMBER:
                rr.bold = True
                rr.font.color.rgb = RGBColor.from_string("B5401F")
    for row in tb.rows:
        row.cells[0].width = Cm(4.2)
        for cc in row.cells[1:]:
            cc.width = Cm(2.2)
        for cc in row.cells:
            cc.paragraphs[0].paragraph_format.space_after = Pt(0)
    para("Each plant's rate per 100,000 m³ injected (the operators' reported volumes, each 2026 report counted for "
         "the month it reports) over its complete record, Pfeilbrunn's from 23:30 to 05:00 scaled to the day and "
         "Lanzenried's from August 2025, times the 2027 plan volume: each quarter as in 2026, the fourth as in 2025. Amber "
         "in bold. Another rate convention on the same record, such as Pfeilbrunn's night count against the volume "
         "injected in those hours or a window from midnight, moves a cell by up to %s per cent and none across the line."
         % TOL, size=8, color="56606A", after=10)

    para("Counting as in section 3 of the 2024 scheme: epicentres within 5 km of the open hole of an injection well, ML "
         "of our reviewed catalogue, calendar quarters in local time. The 2026 measures for the annex to the operating "
         "plans are in the accompanying workbook.", size=9)
    para("G. Atzler", after=0)
    para("Head of induced seismicity", size=9, color="56606A")

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = footer.add_run("LDG-SD-AO-2026-07   page ")
    fr.font.size = Pt(7.5)
    run = footer.add_run()
    run.font.size = Pt(7.5)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    doc.core_properties.title = "Traffic-light outlook 2027"
    doc.save(path)
    PP.normalize_ooxml(path, LEAD, LEAD, "2026-10-12T11:02:00Z", "2026-10-12T15:41:00Z", title="Traffic-light outlook 2027")


# ------------------------------------------------------------------------------------------------- the annex
def write_xlsx(path, c):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    a = c["asks"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Annex 2026"
    ws["A1"] = "Annex to the 2027 operating plans: traffic-light measures, January to September 2026"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Landesdienst für Geologie, Seismologischer Dienst, for the %s under section 7.2 of the scheme. " \
               "Records to 30 September 2026; quarters in local time." % W.AUTHORITY
    ws["A2"].font = Font(italic=True, size=9, color="56606A")
    head = ["Plant", "Quarter", "Injected volume (m³)", "Earthquakes located by the automatic system "
            "(ML 1.0 or more, within 5 km)", "Hours with fewer than four stations recording"]
    fill = PatternFill("solid", fgColor="E6EAEE")
    for j, h in enumerate(head, 1):
        cell = ws.cell(row=4, column=j, value=h)
        cell.font = Font(bold=True, size=10)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left" if j < 3 else "right")
    ws.row_dimensions[4].height = 46
    row = 5
    hair = Border(bottom=Side(style="hair", color="C8CDD2"))
    for p in W.PLANTS:
        for y, q in A.QS:
            x = a[p][(y, q)]
            vals = [W.NAME[p], "Q%d 2026" % q, x["volume_m3"], x["automatic"], x["hours_under_four"]]
            for j, v in enumerate(vals, 1):
                cell = ws.cell(row=row, column=j, value=v)
                cell.border = hair
                if j == 3:
                    cell.number_format = "#,##0"
                elif j == 4:
                    cell.number_format = "0"
                elif j == 5:
                    cell.number_format = "0.0"
            row += 1
    for col, wd in zip("ABCDE", (16, 10, 18, 26, 22)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "C5"
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws = wb.create_sheet("Notes")
    notes = [
        "Injected volume: the operators' monthly reports of metered volume (scheme 6.2), summed over each plant's "
        "injection wells for the three months of the quarter, rounded to the nearest 100 m³. Since January 2026 the "
        "reports come through the authority's portal, whose export files each report under the month it was submitted "
        "in, the month after the month it reports; each is counted here for the month it reports. Moorstetten's "
        "second injection well still reports under its drilling name MST-Th4a (MST-GT4 in the well register).",
        "Automatic system: origins authored by scautoloc in the event database's origin table. The system computes "
        "several origins for most events as picks come in; each event is counted once, at the last origin the system "
        "computed, by that origin's epicentre and ML. The analysts' review origins in the same table are not the "
        "automatic system's.",
        "Hours: a station records while its vertical channel is in an open epoch of the station inventory and outside "
        "the gaps of the availability log (scheme 5.2). On a digitiser exchange day the old epoch ends and the new one "
        "starts some hours later; the station is not recording in between. Networks as in the plant register: "
        "Lanzenried with LZR5 in place of LZR4.",
        "Counting zone for the automatic system: epicentre within 5 km of the open hole of an injection well, as in "
        "scheme 2.1. Quarters are calendar quarters in local time; data times are UTC.",
    ]
    ws["A1"] = "Notes"
    ws["A1"].font = Font(bold=True)
    for i, t in enumerate(notes, 2):
        cell = ws.cell(row=i, column=1, value="%d. %s" % (i - 1, t))
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 58
    ws.column_dimensions["A"].width = 118
    wb.active = 0
    wb.save(path)
    PP.normalize_ooxml(path, LEAD, LEAD, "2026-10-12T09:15:00Z", "2026-10-12T15:38:00Z", excel=True)


def main(target, golden, figures_path):
    os.makedirs(golden, exist_ok=True)
    c = compute(target)
    fig = chart_png(c)
    write_docx(os.path.join(golden, DOCX), c, fig)
    write_xlsx(os.path.join(golden, XLSX), c)
    with open(figures_path, "w") as fh:
        json.dump(figures(c), fh, indent=1, ensure_ascii=False)
        fh.write("\n")
    return c
