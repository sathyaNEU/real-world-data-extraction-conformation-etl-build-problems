"""The pack's documents: the quarter's screen of spontaneous reports with its line listing (the workbook the unit
works from), SOP-PV-14 (the signal refinement procedure), the registry data manager's notes on the extract, the
marketing authorisation holder's PSUR summary for tasorexant, and the minutes of the July committee meeting."""
import datetime as dt
import math

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                           # noqa: E402
from reportlab.lib.pagesizes import A4                                     # noqa: E402
from reportlab.lib.styles import ParagraphStyle                            # noqa: E402
from reportlab.lib.units import mm                                         # noqa: E402
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import postprocess as PP
import world as W

SCREEN = "pv_signal_screen_2026Q3.xlsx"
SOP = "SOP-PV-14_signal_refinement.pdf"
NOTES = "registry_extract_notes.md"
PSUR = "psur_summary_tasorexant_2026.pdf"
MINUTES = "committee_minutes_2026-07-09.md"

INK = colors.HexColor("#1f2a36")
RULE = colors.HexColor("#8a96a3")
GREY = colors.HexColor("#4b5563")
BODY = ParagraphStyle("body", fontName="Helvetica", fontSize=9.3, leading=12.4, textColor=INK, spaceAfter=4)
CLAUSE = ParagraphStyle("clause", parent=BODY, leftIndent=24, firstLineIndent=-24)
H1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=14.5, leading=18, textColor=INK, spaceAfter=2)
H2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=10.4, leading=13.5, textColor=INK, spaceBefore=7,
                    spaceAfter=3)
SMALL = ParagraphStyle("small", parent=BODY, fontSize=7.8, leading=10, textColor=GREY)
CELL = ParagraphStyle("cell", parent=BODY, fontSize=8.2, leading=10.2, spaceAfter=0)
CELLB = ParagraphStyle("cellb", parent=CELL, fontName="Helvetica-Bold")


def _frame(left, right, foot):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.4)
        canvas.setFillColor(GREY)
        canvas.drawString(18 * mm, A4[1] - 12 * mm, left)
        canvas.drawRightString(A4[0] - 18 * mm, A4[1] - 12 * mm, right)
        canvas.setStrokeColor(RULE)
        canvas.setLineWidth(0.4)
        canvas.line(18 * mm, A4[1] - 13.6 * mm, A4[0] - 18 * mm, A4[1] - 13.6 * mm)
        canvas.drawString(18 * mm, 11 * mm, foot)
        canvas.drawRightString(A4[0] - 18 * mm, 11 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _doc(path, title, author, subject, creator):
    return SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=20 * mm,
                             bottomMargin=18 * mm, title=title, author=author, subject=subject, creator=creator,
                             producer=creator)


def _table(rows, widths, header=True, right=()):
    data = [[Paragraph(str(c), CELLB if (header and i == 0) else CELL) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, hAlign="LEFT", repeatRows=1 if header else 0)
    st = [("LINEBELOW", (0, 0), (-1, 0), 0.6, RULE), ("LINEBELOW", (0, -1), (-1, -1), 0.4, RULE),
          ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 2.0),
          ("BOTTOMPADDING", (0, 0), (-1, -1), 2.0), ("LEFTPADDING", (0, 0), (-1, -1), 3)]
    for c in right:
        st.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
    t.setStyle(TableStyle(st))
    return t


# ------------------------------------------------------------------ SOP-PV-14

SOP_SECTIONS = [
    ("1", "Purpose", [
        ("1.1", "This procedure sets how the Pharmacovigilance Unit refines, in the national registries, a signal "
                "validated from the quarterly screen of spontaneous reports, and what it reports to the Medicines "
                "Safety Committee before the committee decides whether to open a formal assessment."),
        ("1.2", "The committee asks one question of each signal: does starting the medicine raise the risk of the "
                "event in the month after the start? A signal is supported when it does."),
    ]),
    ("2", "Responsibilities", [
        ("2.1", "The signal assessor validates signals from the screen and owns this procedure. The registry data "
                "manager orders the registry extract from the national health data service and documents it. The "
                "pharmacoepidemiologist carries out the refinement and writes the note to the committee; the medical "
                "assessor reviews the case definitions against the reports. The head of unit approves the note."),
    ]),
    ("3", "Data", [
        ("3.1", "Refinement uses the registry extract delivered for the quarter: the prescription registry, the "
                "patient registry and the civil registration system, linked on the pseudonymised personal number. The "
                "registry data manager's notes describe each file, its fields and its coverage."),
        ("3.2", "Spontaneous reports are not used to estimate risk. They are used to frame the event and to check "
                "the case definition against what reporters describe."),
    ]),
    ("4", "Events", [
        ("4.1", "Each event has a case definition in annex A. Diagnoses are the patient registry's ICD-10 codes as "
                "the registry holds them, with the D prefix and without the dot; a code includes its subcodes."),
        ("4.2", "Contacts linked by transfer between departments make one admission and count as one event. An "
                "event is dated by the start of its admission."),
    ]),
    ("5", "New users", [
        ("5.1", "A new user is a person dispensed the substance who had no dispensing of it in the 365 days before "
                "and who was resident in Denmark for the whole of those 365 days. The day of that dispensing is the "
                "start (day 0)."),
        ("5.2", "A substance is identified by every ATC code it has carried, so a change of code does not make a "
                "new user."),
        ("5.3", "A dispensing returned to the pharmacy does not count as a dispensing."),
        ("5.4", "A person can be a new user more than once, each time after 365 days without the substance."),
    ]),
    ("6", "The month after the start", [
        ("6.1", "The month after the start is days 1 to 30 after the start."),
    ]),
    ("7", "Observation", [
        ("7.1", "A person's observation ends at death, at emigration or on the last day of data in the extract, "
                "whichever comes first."),
    ]),
    ("8", "The note to the committee", [
        ("8.1", "The note names the signal the unit recommends for a formal assessment and gives, for that "
                "medicine, the rate ratio for the month after the start with its 95 per cent confidence interval, the "
                "admissions for the event in that month among people who started the medicine in the last full "
                "calendar year of the extract, and how many of those admissions the medicine caused."),
        ("8.2", "Rate ratios and their confidence limits are given to two decimals. Admissions caused by the "
                "medicine are the admissions in the month times (ratio - 1) / ratio, rounded to the nearest whole "
                "admission."),
        ("8.3", "The note goes to the committee secretary by the Monday before the meeting."),
    ]),
]

ANNEX_A = [
    ["Event", "Case definition", "Codes"],
    ["Fall", "An emergency contact or an inpatient admission whose action diagnosis is an injury and which carries a "
             "fall code among its diagnoses", "Injury DS00 to DT14; fall DW00 to DW19"],
    ["Acute kidney injury", "An inpatient admission with acute kidney injury as action diagnosis", "DN17"],
    ["Acute pancreatitis", "An inpatient admission with acute pancreatitis as action diagnosis", "DK85"],
]


def write_sop(path):
    title = "SOP-PV-14 Signal refinement in the national registries"
    doc = _doc(path, title, W.PEOPLE["assessor"], "Standard operating procedure, revision 3", "Pharmacovigilance Unit")
    s = [Paragraph("The Agency, Pharmacovigilance Unit", SMALL),
         Paragraph(title, H1), Spacer(1, 4),
         _table([["Document", "Revision", "Effective", "Owner", "Approved"],
                 ["SOP-PV-14", "3", "1 June 2026", "%s, signal assessor" % W.PEOPLE["assessor"],
                  "%s, head of unit, 20 May 2026" % W.PEOPLE["head"]]],
                [22 * mm, 16 * mm, 22 * mm, 50 * mm, 64 * mm]),
         Spacer(1, 6)]
    for num, head, clauses in SOP_SECTIONS:
        block = [Paragraph("%s  %s" % (num, head), H2)]
        for cl, text in clauses:
            block.append(Paragraph("<b>%s</b>&nbsp;&nbsp;%s" % (cl, text), CLAUSE))
        s.append(KeepTogether(block))
    s += [Spacer(1, 6), KeepTogether([Paragraph("Annex A  Case definitions", H2),
                                      _table(ANNEX_A, [34 * mm, 92 * mm, 48 * mm])])]
    s += [Spacer(1, 6), Paragraph("Revision history", H2),
          _table([["Revision", "Date", "Change"],
                  ["3", "20 May 2026", "Sections 4 and 5 rewritten; note to the committee (section 8)"],
                  ["2", "14 March 2022", "Patient registry contacts from 2019 (section 4)"],
                  ["1", "2 September 2019", "First issue"]], [20 * mm, 28 * mm, 120 * mm]),
          Spacer(1, 6),
          Paragraph("Printed copies are uncontrolled. The current revision is in the unit's document system.", SMALL)]
    foot = "SOP-PV-14 revision 3, effective 1 June 2026"
    doc.build(s, onFirstPage=_frame("The Agency", "SOP-PV-14", foot), onLaterPages=_frame("The Agency", "SOP-PV-14", foot))
    PP.normalize_pdf(path, b"D:20260520141500+02'00'")


# ------------------------------------------------------------------ the extract notes

def write_notes(path, counts):
    text = """# Registry extract PV-2026-117: notes on the files

Pharmacovigilance Unit. %(data)s, registry data manager. Delivered by the national health data service on 4 September
2026, data to 30 June 2026. Ordered on 14 July 2026, after the Q2 screen, for the three products whose pairs the unit
was following (tasorexant, corinurad and velaglutide), so that refinement can start as soon as a pair is validated.

## Who is in the extract

- Everyone dispensed tasorexant, corinurad or velaglutide up to 30 June 2026.
- Everyone the data service found starting zopiclone, allopurinol or orlistat between 1 January 2022 and 30 June 2026
  after a year without it (their selection, on the raw dispensings).
- The matched comparators drawn for the first group (`matched_comparators.csv`).

For each of them the extract holds every dispensing of the six substances and of NSAIDs (ATC M01A) from 1 January 2018,
every hospital contact from 1 January 2018 with its diagnoses, and their civil registration record. Personal numbers are
pseudonymised the same way in every file.

## Files

| File | Rows | What one row is |
|---|---|---|
| `%(rx)s` | %(n_rx)s | one dispensing (or return) at a pharmacy |
| `%(c)s` | %(n_c)s | one hospital record (see the 2019 change below) |
| `%(dx)s` | %(n_dx)s | one diagnosis on a hospital record |
| `%(civ)s` | %(n_civ)s | one person |
| `%(mt)s` | %(n_mt)s | one person in one matched set |
| `%(atc)s` | %(n_atc)s | one ATC code |
| `%(pop)s` | %(n_pop)s | one year, age group and sex, nationally |

## Prescription registry

| Field | Meaning |
|---|---|
| `pnr` | pseudonymised personal number |
| `eksd` | dispensing date |
| `atc` | ATC code of the product, as valid on the dispensing date (codes and their validity: `%(atc)s`) |
| `vnr` | package number (varenummer) |
| `apk` | number of packages; a package returned to the pharmacy is a row with a negative number of packages on the day of return, against the same package number |
| `packsize` | units in the package |
| `strnum`, `strunit` | strength per unit |
| `volume`, `voltypetxt` | defined daily doses in one package |

## Patient registry

The registry changed structure on 4 February 2019, and the column `kilde` says which structure a record has.

- **LPR2** (to 3 February 2019). An inpatient stay (`pattype` 0) or an emergency room visit (`pattype` 3) is one record per
  department, with the admission date and hour in `starttid` and the discharge date in `sluttid`. Outpatient care
  (`pattype` 2) is one record per course: `starttid` is the first visit and `sluttid` the last, and the course's
  diagnoses are on that one record. A course still running on 4 February 2019 ends at its last LPR2 visit; its later
  visits are LPR3 contacts.
- **LPR3** (from 4 February 2019). Every stay at a department is a contact, with `starttid` and `sluttid` to the minute,
  `prioritet` ATA1 (acute) or ATA3 (planned) and `kontakttype` ALCA00 (physical attendance) or ALCA10 (virtual).
  `sluttid` is empty for a contact still open when the extract was cut.

A patient moved between departments during a stay has a record at each department. The data service treats LPR3
contacts as one admission when one starts within four hours of the end of the previous one, and LPR2 records as one
admission when one starts on the date the previous one ends. An inpatient admission is an admission with physical
attendance of 12 hours or more (LPR3) or a `pattype` 0 record (LPR2); an emergency contact is an acute physical
attendance (LPR3) or a `pattype` 3 record (LPR2).

| Field | Meaning |
|---|---|
| `kontakt_id` | the record's key; links the diagnoses |
| `afdeling` | department: SHAK code (LPR2) or SOR code (LPR3) |
| `speciale` | the department's specialty, from the department register |
| `diagnosekode` | ICD-10 code as the registry holds it: a D prefix and no dot (DE119 is E11.9) |
| `diagnosetype` | A action diagnosis (the main reason for the contact), B secondary diagnosis |

## Civil registration

`foedselsaar` birth year; `koen` M or K; `doedsdato` date of death. `udvandringsdato` and `indvandringsdato` are the
person's most recent emigration and immigration up to the extract date: someone who left and came back has both, with
the immigration after the emigration. Blank means none on record.

## Matched comparators

The Agency's standard matched extract. For each person's first dispensing of tasorexant, corinurad or velaglutide (the
index date, `indeksdato`) the data service drew two people of the same birth year and sex who were alive and resident on
that day and had not been dispensed the substance before it. `saet_id` groups the person dispensed (`rolle` bruger) with
the two drawn (`rolle` sammenligning). The draw was made on the raw dispensings.

## Population and national admissions

Population on 1 January by five-year age group and sex, aged 20 and over, with admissions counted nationally by the data
service under the case definitions of SOP-PV-14 annex A, 2022 to 2025.
""" % dict(data=W.PEOPLE["data"], rx="prescription_registry_2018-2026.parquet",
           c="patient_registry_contacts_2018-2026.parquet", dx="patient_registry_diagnoses_2018-2026.parquet",
           civ="civil_registration_extract.csv", mt="matched_comparators.csv", atc="atc_ddd_index_2026.csv",
           pop="population_and_admissions_2022-2025.csv",
           n_rx="{:,}".format(counts["rx"]), n_c="{:,}".format(counts["contacts"]), n_dx="{:,}".format(counts["dx"]),
           n_civ="{:,}".format(counts["civil"]), n_mt="{:,}".format(counts["matched"]), n_atc=counts["atc"],
           n_pop=counts["pop"])
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(text)


# ------------------------------------------------------------------ the screen workbook

SCREEN_ROWS = [
    # substance, product, PT, Q3 cases, cumulative cases, reports with substance, reports with PT, status
    ("allopurinol", "", "Drug reaction with eosinophilia and systemic symptoms", 3, 141, 2218, 1904, "Labelled"),
    ("corinurad", "Uralexa", "Acute kidney injury", None, 24, 286, 6120, "Validated signal, to refinement"),
    ("corinurad", "Uralexa", "Nephrolithiasis", 3, 12, 286, 1630, "Labelled"),
    ("corinurad", "Uralexa", "Gout", 5, 27, 286, 2391, "Labelled"),
    ("tasorexant", "Somnevo", "Fall", None, 128, 412, 18340, "Validated signal, to refinement"),
    ("tasorexant", "Somnevo", "Nightmare", 4, 22, 412, 2867, "Labelled"),
    ("tasorexant", "Somnevo", "Somnolence", 11, 64, 412, 12405, "Labelled"),
    ("velaglutide", "Gravera", "Cholelithiasis", 9, 31, 1904, 2512, "Labelled"),
    ("velaglutide", "Gravera", "Nausea", 48, 312, 1904, 28733, "Labelled"),
    ("velaglutide", "Gravera", "Pancreatitis acute", None, 57, 1904, 3880, "Validated signal, to refinement"),
    ("zopiclone", "", "Dysgeusia", 6, 488, 3104, 3011, "Labelled"),
    ("apixaban", "", "Gastrointestinal haemorrhage", 23, 871, 9126, 6214, "Labelled"),
    ("dapagliflozin", "", "Diabetic ketoacidosis", 7, 314, 5479, 1961, "Labelled"),
    ("dupilumab", "", "Conjunctivitis", 12, 643, 4108, 3307, "Labelled"),
    ("finasteride", "", "Depression", 4, 58, 1283, 9034, "Closed: not confirmed (Q4 2025)"),
    ("levetiracetam", "", "Aggression", 5, 197, 3315, 2243, "Labelled"),
    ("metformin", "", "Lactic acidosis", 3, 152, 7806, 401, "Labelled"),
    ("semaglutide", "", "Cholelithiasis", 26, 417, 15209, 2512, "Labelled"),
    ("varenicline", "", "Abnormal dreams", 4, 236, 2953, 1118, "Labelled"),
]
DATABASE_REPORTS = 812446

CONCOMITANT_POOL = ["amlodipine", "atorvastatin", "metformin", "ramipril", "bisoprolol", "furosemide", "levothyroxine",
                    "omeprazole", "paracetamol", "simvastatin", "losartan", "apixaban", "sertraline", "alendronic acid",
                    "colchicine", "bendroflumethiazide", "candesartan", "pantoprazole"]


def screen_cases():
    """The quarter's reports for the three signals: one row per case version."""
    r = W.rng("reports")
    spec = {
        "tasorexant": dict(n=34, brand="Somnevo", pts=[["Fall"], ["Fall", "Hip fracture"], ["Fall", "Head injury"],
                                                       ["Fall", "Contusion"], ["Fall", "Wrist fracture"]],
                           age=(81, 7, 62, 97), female=0.7),
        "corinurad": dict(n=13, brand="Uralexa", pts=[["Acute kidney injury"], ["Acute kidney injury",
                                                                                "Blood creatinine increased"],
                                                      ["Acute kidney injury", "Gout"]],
                          age=(66, 9, 44, 86), female=0.15),
        "velaglutide": dict(n=21, brand="Gravera", pts=[["Pancreatitis acute"], ["Pancreatitis acute", "Abdominal pain"],
                                                        ["Pancreatitis acute", "Lipase increased"],
                                                        ["Pancreatitis acute", "Vomiting"]],
                            age=(52, 10, 27, 74), female=0.62),
    }
    rows = []
    q0, q1 = W.day(W.Q3[0]), W.day(W.Q3[1])
    rc = W.rng("reports", "received")
    # a case number is given when a case is first received, so numbers follow receipt across products
    first = [(int(rc.integers(q0, q1 - 20)), sub, k) for sub, sp in spec.items() for k in range(sp["n"])]
    serial_of, serial = {}, 4180
    for rec_, sub, k in sorted(first):
        serial += int(rc.integers(3, 40))
        serial_of[(sub, k)] = (serial, rec_)
    for sub, sp in spec.items():
        for k in range(sp["n"]):
            serial, rec = serial_of[(sub, k)]
            case = "DK-PV-26-%05d" % serial
            age = int(np.clip(round(r.normal(sp["age"][0], sp["age"][1])), sp["age"][2], sp["age"][3]))
            sex = "F" if r.random() < sp["female"] else "M"
            reporter = str(r.choice(["Physician", "Pharmacist", "Consumer", "Nurse", "Other health professional"],
                                    p=[0.42, 0.18, 0.22, 0.12, 0.06]))
            pts = sp["pts"][int(r.integers(0, len(sp["pts"])))]
            suspect = "%s (%s)" % (sub, sp["brand"])
            if sub == "tasorexant" and r.random() < 0.12:
                suspect += "; zopiclone"
            con = "; ".join(sorted(str(x) for x in r.choice(CONCOMITANT_POOL, size=int(r.integers(0, 4)), replace=False)))
            started = rec - int(r.integers(5, 700))
            onset = started + int(r.integers(1, max(2, rec - started)))
            s_txt = _partial(r, started)
            o_txt = _partial(r, onset)
            serious = "Yes" if r.random() < 0.8 else "No"
            outcome = str(r.choice(["Recovered", "Recovering", "Not recovered", "Unknown", "Recovered with sequelae"],
                                   p=[0.35, 0.3, 0.1, 0.2, 0.05]))
            versions = 1 + (1 if r.random() < 0.3 else 0) + (1 if r.random() < 0.08 else 0)
            for v in range(1, versions + 1):
                vrec = rec + (0 if v == 1 else int(r.integers(3, 18)) * (v - 1))
                vrec = min(vrec, q1)
                out = outcome if v == versions else "Unknown"
                rows.append([case, v, W.date_of(vrec), "Spontaneous", reporter, age if (v > 1 or r.random() < 0.9) else None,
                             sex, suspect, con, "; ".join(pts), serious, out, s_txt, o_txt if v > 1 or r.random() < 0.7
                             else ""])
    # reports where the medicine was taken but not suspected
    used = {v[0] for v in serial_of.values()}
    for sub, pt, sus in (("tasorexant", "Fall", "zopiclone"), ("tasorexant", "Fall", "mirtazapine"),
                         ("corinurad", "Acute kidney injury", "ibuprofen"), ("velaglutide", "Pancreatitis acute",
                                                                              "azathioprine")):
        rec = int(r.integers(q0, q1 - 5))
        prior = [v[0] for v in serial_of.values() if v[1] <= rec]
        serial = (max(prior) if prior else 4180) + 1
        while serial in used:
            serial += 1
        used.add(serial)
        rows.append(["DK-PV-26-%05d" % serial, 1, W.date_of(rec), "Spontaneous", "Physician",
                     int(r.integers(55, 88)), "F" if r.random() < 0.5 else "M", sus,
                     "%s (%s)" % (sub, W.SUBSTANCES[sub]["brand"]), pt, "Yes", "Recovering", "", ""])
    rows.sort(key=lambda x: (x[0], x[1]))
    return rows


def _partial(r, d):
    x = r.random()
    date = W.date_of(d)
    if x < 0.6:
        return date.isoformat()
    if x < 0.85:
        return date.strftime("%Y-%m")
    return ""


def screen_counts(rows):
    latest = {}
    for row in rows:
        latest[row[0]] = row
    out = {}
    for sub in ("tasorexant", "corinurad", "velaglutide"):
        pt = {"tasorexant": "Fall", "corinurad": "Acute kidney injury", "velaglutide": "Pancreatitis acute"}[sub]
        out[sub] = sum(1 for row in latest.values() if sub in row[7] and pt in row[9].split("; "))
    return out


def prr(a, n_sub, n_pt, N):
    b = n_sub - a
    c = n_pt - a
    d = N - a - b - c
    p = (a / (a + b)) / (c / (c + d))
    se = math.sqrt(1 / a - 1 / (a + b) + 1 / c - 1 / (c + d))
    return p, p * math.exp(-1.96 * se)


def write_screen(path):
    rows = screen_cases()
    q3 = screen_counts(rows)
    wb = openpyxl.Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="2F3E4E")
    thin = Side(style="thin", color="B8C0C8")
    ws = wb.active
    ws.title = "About"
    lines = [
        ("The Agency, Pharmacovigilance Unit", True),
        ("Quarterly screen of spontaneous reports, Q3 2026", True),
        ("Reports received 1 July to 30 September 2026. Screen run 6 October 2026 (%s)." % W.PEOPLE["assessor"], False),
        ("", False),
        ("Method: proportional reporting ratio (PRR) of each substance and MedDRA preferred term against every report in "
         "the national reports database. A pair is listed when it has three or more cases and a PRR whose lower 95 per "
         "cent limit is above 1, or when it was listed the quarter before.", False),
        ("Cases are counted once, on their latest version. A case counts for a substance when the substance is suspect.",
         False),
        ("The screen describes how reports are distributed in the reports database. It does not measure how often an "
         "event happens to people who take a medicine; validated signals go to registry refinement under SOP-PV-14.",
         False),
        ("", False),
        ("Sheets: Screen Q3 2026 (the listed pairs); Line listing (case versions received in the quarter for the three "
         "validated signals, and reports in which the medicine was taken but not suspected).", False),
    ]
    for i, (t, b) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = Font(bold=b, size=12 if i == 2 else 10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120
    ws2 = wb.create_sheet("Screen Q3 2026")
    hdr = ["Substance", "Product", "MedDRA preferred term", "Cases Q3 2026", "Cases since launch",
           "Reports with substance", "Reports with term", "Reports in database", "PRR", "PRR lower 95%", "Status"]
    ws2.append(hdr)
    for sub, prod, pt, qn, cum, ns, npt, status in SCREEN_ROWS:
        if qn is None:
            qn = q3[sub]
        p, lo = prr(cum, ns, npt, DATABASE_REPORTS)
        ws2.append([sub, prod, pt, qn, cum, ns, npt, DATABASE_REPORTS, round(p, 2), round(lo, 2), status])
    widths = [14, 10, 44, 13, 16, 20, 17, 18, 8, 13, 30]
    for j, w in enumerate(widths, start=1):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
        c = ws2.cell(row=1, column=j)
        c.font, c.fill = head, fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws2.iter_rows(min_row=2):
        for c in row:
            c.border = Border(bottom=thin)
            if c.column in (4, 5, 6, 7, 8):
                c.number_format = "#,##0"
            if c.column in (9, 10):
                c.number_format = "0.00"
    ws2.freeze_panes = "A2"
    ws3 = wb.create_sheet("Line listing")
    ws3.append(["Line listing, Q3 2026: case versions received 1 July to 30 September 2026"])
    ws3.append(["Suspect: medicines the reporter suspected. Concomitant: medicines taken at the time and not "
                "suspected. Dates are as reported (year and month where the day is not known)."])
    ws3.append([])
    hdr3 = ["Case number", "Version", "Received", "Report type", "Reporter", "Age", "Sex", "Suspect", "Concomitant",
            "Reactions (MedDRA PT)", "Serious", "Outcome", "Suspect started", "Reaction onset"]
    ws3.append(hdr3)
    for row in rows:
        ws3.append(row)
    widths3 = [17, 8, 11, 12, 22, 6, 5, 30, 40, 38, 8, 22, 15, 14]
    for j, w in enumerate(widths3, start=1):
        ws3.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
        c = ws3.cell(row=4, column=j)
        c.font, c.fill = head, fill
    ws3["A1"].font = Font(bold=True, size=11)
    for row in ws3.iter_rows(min_row=5):
        row[2].number_format = "yyyy-mm-dd"
    ws3.freeze_panes = "B5"
    ws3.auto_filter.ref = "A4:N%d" % (4 + len(rows))
    wb.save(path)
    PP.normalize_ooxml(path, W.PEOPLE["assessor"], W.PEOPLE["assessor"], "2026-10-06T07:12:00Z", "2026-10-07T13:41:00Z",
                       title="Q3 2026 screen")
    return q3, rows


# ------------------------------------------------------------------ the PSUR summary (the MAH's document)

def write_psur(path, dk_years):
    """dk_years: the interval's Danish exposure in patient-years, from the tablets dispensed in Denmark."""
    title = "Somnevo (tasorexant) PSUR 4: executive summary"
    doc = _doc(path, title, W.MAH, "Periodic safety update report 4, executive summary and exposure", W.MAH)
    dk = "{:,}".format(int(round(dk_years / 50.0)) * 50)
    s = [Paragraph(W.MAH, SMALL), Paragraph("Somnevo (tasorexant) 25 mg and 50 mg film-coated tablets", H1),
         Paragraph("Periodic safety update report 4: executive summary, with the exposure estimate for Denmark", BODY),
         Spacer(1, 4),
         _table([["Reporting interval", "1 February 2025 to 31 January 2026"],
                 ["Data lock point", "31 January 2026"],
                 ["International birth date", "14 June 2021"],
                 ["Submitted", "8 April 2026 (EU single assessment)"]], [45 * mm, 90 * mm], header=False),
         Spacer(1, 6), Paragraph("Executive summary", H2),
         Paragraph("This is the fourth PSUR for Somnevo. Patient exposure in the interval is estimated at about 151,000 "
                   "patient-years worldwide. No new signal was opened in the interval and no change to the reference "
                   "safety information is proposed. Falls remain an important potential risk in the risk management "
                   "plan. The benefit-risk balance of Somnevo is unchanged.", BODY),
         Paragraph("5.2  Patient exposure from marketing experience", H2),
         Paragraph("Exposure is estimated from sales in tablets, at one tablet a day: about 151,000 patient-years in the "
                   "interval and about 318,000 since the international birth date. At the request of the national "
                   "competent authority the same estimate is given for Denmark: about %s patient-years in the "
                   "interval." % dk, BODY),
         _table([["Age group", "Interval exposure, patient-years", "Share"],
                 ["18 to 64", "46,800", "31%"], ["65 to 74", "39,300", "26%"], ["75 and over", "64,900", "43%"]],
                [40 * mm, 60 * mm, 25 * mm], right=(1, 2)),
         Spacer(1, 4), Paragraph("6  Summary tabulation, postmarketing reports (extract)", H2),
         Paragraph("Spontaneous and solicited reports received by the marketing authorisation holder, cases counted once.",
                   BODY),
         _table([["Preferred term", "Interval", "Cumulative", "Serious, cumulative"],
                 ["Somnolence", "212", "588", "41"], ["Fall", "97", "264", "171"], ["Headache", "88", "247", "6"],
                 ["Nightmare", "61", "198", "3"], ["Dizziness", "57", "176", "18"],
                 ["Sleep paralysis", "19", "52", "2"], ["Hip fracture", "17", "41", "41"],
                 ["Complex sleep behaviour", "6", "15", "9"]],
                [60 * mm, 25 * mm, 28 * mm, 38 * mm], right=(1, 2, 3)),
         Spacer(1, 4), Paragraph("16  Signal and risk evaluation (extract)", H2),
         Paragraph("Falls: the reporting rate in the interval was 64 per 100,000 patient-years, against 71 in the "
                   "previous interval. Most fall reports concern patients aged 75 or over, and in about half the patient "
                   "was taking at least one other medicine with sedative effects.", BODY),
         Paragraph("19  Conclusions and actions", H2),
         Paragraph("Routine pharmacovigilance continues, with falls followed as an important potential risk.", BODY),
         Spacer(1, 6),
         Paragraph("Prepared by the marketing authorisation holder's drug safety department. Confidential: for the "
                   "national competent authority.", SMALL)]
    doc.build(s, onFirstPage=_frame(W.MAH, "Somnevo PSUR 4", "DLP 31 January 2026"),
              onLaterPages=_frame(W.MAH, "Somnevo PSUR 4", "DLP 31 January 2026"))
    PP.normalize_pdf(path, b"D:20260408103000+02'00'")


def psur_exposure(world):
    """Tablets of tasorexant dispensed in Denmark, 1 February 2025 to 31 January 2026, as patient-years at one a day."""
    d = world["disp"]
    lo, hi = W.day(dt.date(2025, 2, 1)), W.day(dt.date(2026, 1, 31))
    m = (d["substance"] == "tasorexant") & (d["day"] >= lo) & (d["day"] <= hi)
    tablets = sum(int(a) * W.PACK[v]["size"] for v, a in zip(d["vnr"][m], d["apk"][m]))
    return tablets / 365.25


# ------------------------------------------------------------------ the July minutes

def write_minutes(path):
    p = W.PEOPLE
    text = """# Medicines Safety Committee: minutes of the meeting of 9 July 2026

The Agency. Meeting held 9 July 2026, 13:00 to 15:10, meeting room 3 and online.

**Present:** the committee chair (head of the Medicines Safety Division), the two external clinical members,
%(head)s (head of the Pharmacovigilance Unit), %(med)s (medical assessor), %(epi)s (pharmacoepidemiologist),
%(ass)s (signal assessor). **Secretary:** %(sec)s.

## 1. Minutes of 16 April 2026

Approved without changes.

## 2. Signals from the Q2 2026 screen

**2.1 Hyponatraemia with velnoxetine.** The unit's registry refinement did not support the signal. The committee closed
it; the product information already lists hyponatraemia as uncommon. Action: none.

**2.2 Angioedema with brelostat.** The refinement supported the signal. The committee opened a formal assessment with
%(med)s as rapporteur; the assessment report is due at the October meeting. Action: %(med)s.

## 3. Tasorexant (Somnevo)

The unit reported that falls with tasorexant were again among the pairs on the Q2 screen, below the threshold for a
validated signal (18 cases in the quarter). The committee asked to be told if the pair is validated in Q3. The marketing
authorisation holder's PSUR (data lock 31 January 2026) is with the unit. Action: %(ass)s.

## 4. SOP-PV-14

The committee noted revision 3 of SOP-PV-14 (signal refinement in the national registries), in force from 1 June 2026.

## 5. Arrangements for October

The Q3 screen will run in the first week of October. Refinement notes on the Q3 signals reach the secretary by Monday
19 October for the meeting on Thursday 22 October 2026. The committee has room for one new formal assessment at that
meeting, because the brelostat assessment is still running.

## 6. Any other business

None.

%(sec)s, secretary, 13 July 2026
""" % dict(head=p["head"], med=p["medical"], epi=p["epi"], ass=p["assessor"], sec=p["secretary"])
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(text)
