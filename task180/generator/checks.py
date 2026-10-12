"""The generator's assertions on the built pack and goldens: the ladder (every rung's answer by name), the semantic
forks SF1 to SF10 (each reading's call), the margins, the three asks (golden against the simulation's truth, every
cracker-path figure off), the separation line, the input gates and the pack's hygiene. Any failure stops the build."""
import datetime as dt
import json
import os
import re
import zipfile

import numpy as np
import pandas as pd

import asks as A
import geo as G
import models as M
import records as R
import sim as S
import world as W

N = [0]
GOLD = [("PFB", 1), ("PFB", 4)]


def ok(cond, label):
    N[0] += 1
    if not cond:
        raise AssertionError("check failed: %s" % label)


def main_call(target, world):
    d = M.load(target)
    c = M.counted(d)
    g = M.golden(d, c)
    vol, _ = M.volumes(d)
    rq = M.report_quarters(d["months"])
    plan = M.plan_reports(d)
    # the stated tolerance: every rate convention on the whole record keeps every cell within 8 per cent and the call
    for name, kw in (("telemetry volumes", dict(source="telemetry")), ("night count over the night's volume", dict(night_volume=True)),
                     ("window 00:00 to 05:00", dict(window=(0.0, 5.0))), ("window 23:00 to 05:00", dict(window=(23.0, 5.0))),
                     ("window 00:00 to 04:00", dict(window=(0.0, 4.0))), ("window 00:30 to 04:30", dict(window=(0.5, 4.5))),
                     ("record from 2024", dict(years=(2024, 2025, 2026))),
                     ("Lanzenried from the switch day (telemetry)", dict(source="telemetry", lzr_from=True))):
        f = M.golden(d, c, **kw)
        ok(M.over(f) == GOLD, "convention %s keeps the call" % name)
        dev = max(abs(f[p][q] / g[p][q] - 1) for p in W.PLANTS for q in (1, 2, 3, 4))
        ok(dev <= 0.08, "convention %s keeps every cell within 8 per cent (%.3f)" % (name, dev))
    ok(M.over(g) == GOLD, "golden names Pfeilbrunn in Q1 and Q4 only: %s" % M.over(g))
    r0, r1, r2 = M.rung0(c), M.rung1(c), M.rung2(c, d)
    ok(M.over(r0) == [], "rung 0 names no plant (SF5: the count the 2026 networks located, carried forward)")
    ok(max(r0[p][q] for p in W.PLANTS for q in (1, 2, 3, 4)) <= 0.85 * W.AMBER, "rung 0 nearest cell at most 0.85 of amber")
    ok(M.over(r1) == [("LZR", 1), ("LZR", 4)], "rung 1 names Lanzenried in Q1 and Q4: %s" % M.over(r1))
    ok(M.over(r2) == [], "rung 2 names no plant")
    for q in (1, 4):
        ok(any(r0[p][q] >= r0["PFB"][q] for p in W.PLANTS if p != "PFB"), "Pfeilbrunn not alone at the top of rung 0 in Q%d" % q)
        ok(r0["PFB"][q] <= 0.70 * W.AMBER, "Pfeilbrunn's rung-0 Q%d at most 0.70 of amber" % q)
        ok(g["PFB"][q] >= 1.5 * W.AMBER, "Pfeilbrunn's golden Q%d at least 1.5 times amber" % q)
        ok(g["PFB"][q] / r0["PFB"][q] >= 2.2, "dominance: the decisive move multiplies Pfeilbrunn's Q%d by 2.2 or more" % q)
    for p in W.PLANTS:
        for q in (1, 2, 3, 4):
            if (p, q) not in GOLD:
                ok(g[p][q] <= 0.75 * W.AMBER, "golden %s Q%d at most 0.75 of amber" % (p, q))
            ok(abs(g[p][q] - W.AMBER) > 3, "SF7 no golden cell within 3 of the line (%s Q%d)" % (p, q))
    # the grid: growth, Lanzenried's network change, completeness by hour; eight cells, each named
    t0 = M.table(c)

    def growth(p, q):
        y1, y0 = ((2026, 2025) if q < 4 else (2025, 2024))
        return t0[p][(y1, q)] / t0[p][(y0, q)] if t0[p][(y0, q)] else 1.0
    for gr in (0, 1):
        for net in (0, 1):
            for hr in (0, 1):
                f = M.forecast_from_rates(M.rates(c, d, quiet_plants=("PFB",) if hr else (), lzr_from=bool(net)), plan)
                if gr:
                    f = {p: {q: f[p][q] * growth(p, q) for q in f[p]} for p in f}
                ov = M.over(f)
                ok((("PFB", 1) in ov and ("PFB", 4) in ov) == bool(hr) and not any(x[0] == "PFB" and x[1] in (2, 3) for x in ov),
                   "grid: Pfeilbrunn over in Q1 and Q4 exactly when completeness by hour (g%d n%d h%d)" % (gr, net, hr))
                ok(any(x[0] == "LZR" for x in ov) == bool(gr and net), "grid: Lanzenried over only with growth on its new network (g%d n%d h%d)" % (gr, net, hr))
                ok(max(f[p][q] for p in ("MOS", "WDR", "OTM") for q in (1, 2, 3, 4)) <= 0.8 * W.AMBER,
                   "grid: Moorstetten, Wendelsried and Ottenmoos under 0.8 of amber (g%d n%d h%d)" % (gr, net, hr))
    # SF8 windows; SF10 quiet windows; the quiet-hour estimator on every plant
    for yrs in [(2026,), (2025, 2026), (2024, 2025, 2026), (2023, 2024, 2025, 2026)]:
        ok(M.over(M.forecast_from_rates(M.rates(c, d, years=yrs), plan)) == GOLD, "SF8 window %s" % (yrs,))
        ok(M.over(M.forecast_from_rates(M.rates(c, d, years=yrs, quiet_plants=tuple(W.PLANTS)), plan)) == GOLD,
           "quiet-hour estimator on every plant, window %s" % (yrs,))
    for win in [(23.5, 5.0), (0.0, 4.0), (0.5, 4.5), (23.0, 5.0), (22.0, 6.0), (1.0, 5.0)]:
        ok(M.over(M.forecast_from_rates(M.rates(c, d, window=win), plan)) == GOLD, "SF10 quiet window %s" % (win,))
    # single quarters: Pfeilbrunn's quiet-hour estimate over amber in every winter quarter, under in every summer one
    for (y, q), n in M.table(c[M.quiet(c["h"].to_numpy())])["PFB"].items():
        est = n / M.SHARE
        ok((est > W.AMBER) == (q in (1, 4)), "Pfeilbrunn quiet-hour estimate %d Q%d on its side of amber (%.0f)" % (y, q, est))
    # SF2 the open hole against the wellhead and the open hole's midpoint
    wells = d["wells"]
    for variant in ("wellhead", "midpoint"):
        hm = {}
        for x in wells[wells["function"] == "injection"].itertuples():
            if variant == "wellhead":
                pt = (x.wellhead_lat, x.wellhead_lon, 0)
            else:
                pt = ((x.oh_top_lat + x.oh_bottom_lat) / 2, (x.oh_top_lon + x.oh_bottom_lon) / 2, 0)
            hm.setdefault(x.plant, []).append((pt, pt))
        cc = d["cat"][(d["cat"]["event_type"] == "earthquake") & (d["cat"]["ml"] >= 1.0)].copy()
        cc["plant"], cc["dist"] = M.zone(cc["latitude"], cc["longitude"], hm)
        cc = cc[cc["plant"] != ""].copy()
        cc["y"], cc["q"], cc["h"] = M.local_parts(cc["t"].to_numpy())
        ok(M.over(M.golden(d, cc)) == GOLD, "SF2 %s reading gives the same call" % variant)
    # SF3 UTC quarters
    cu = c.copy()
    ts = pd.to_datetime(cu["t"].to_numpy(), unit="s", utc=True)
    cu["y"], cu["q"] = ts.year.to_numpy(), ts.quarter.to_numpy()
    ok(M.over(M.golden(d, cu)) == GOLD, "SF3 UTC quarters give the same call")
    # SF1 no explosion within 5 km of an injection open hole; SF9 the median of the forecast
    ex = d["cat"][d["cat"]["event_type"] == "explosion"]
    pl, dist = M.zone(ex["latitude"], ex["longitude"], M.holes(wells))
    ok((pl != "").sum() == 0 and dist.min() > 8.0, "SF1 explosions all beyond 8 km of every open hole (%.1f)" % dist.min())
    from scipy.stats import poisson
    ok(sorted((p, q) for p in W.PLANTS for q in (1, 2, 3, 4) if poisson.median(g[p][q]) > W.AMBER) == GOLD,
       "SF9 the median count gives the same call")
    # SF4 the reviewed catalogue is what the scheme counts; the federal catalogue cannot count ML 1.0
    fed = pd.read_csv(os.path.join(target, R.FEDERAL))
    ok(fed["mag"].min() >= 2.0, "SF4 the federal catalogue starts at magnitude 2.0")
    # SF6 the plans: 2027 quarter volumes are the record's 2026 (and Q4 2025) volumes
    for p in W.PLANTS:
        ok(plan[p][1] == rq[p][(2026, 1)] and plan[p][4] == rq[p][(2025, 4)], "SF6 plan volumes for %s" % p)
        for k in W.record_quarters():
            ok(abs(rq[p][k] / vol[p][k] - 1) < 0.01, "reported volume ties to the telemetry within 1 per cent (%s %s)" % (p, k))
    # rung 2's premise: maximum curvature reads every plant's record complete at or below ML 0.8; Lanzenried's old
    # network not; Pfeilbrunn's hourly profile carries the decisive signal, the rural ones are flat
    allc = d["cat"][d["cat"]["event_type"] == "earthquake"].copy()
    allc["plant"], _ = M.zone(allc["latitude"], allc["longitude"], M.holes(wells))
    allc = allc[allc["plant"] != ""].copy()
    allc["y"], allc["q"], allc["h"] = M.local_parts(allc["t"].to_numpy())
    for p in W.PLANTS:
        e = allc[allc["plant"] == p]
        if p == "LZR":
            ok(M.maxc(e[e["t"] < d["switch"]]["ml"]) >= 1.2, "Lanzenried's old network complete only above ML 1.2")
            e = e[e["t"] >= d["switch"]]
        mc = M.maxc(e["ml"])
        ok(mc <= 0.8, "maximum curvature reads %s complete at ML %.1f" % (p, mc))
        cp = c[c["plant"] == p]
        if p == "LZR":
            cp = cp[cp["t"] >= d["switch"]]
        share = M.quiet(cp["h"].to_numpy()).mean()
        if p == "PFB":
            ok(share >= 0.5, "Pfeilbrunn's quiet hours hold half or more of its located earthquakes (%.2f)" % share)
        else:
            ok(abs(share - M.SHARE) <= 0.07, "%s's quiet-hour share near 5.5/24 (%.2f)" % (p, share))
    b_pfb, _ = M.bvalue(allc[allc["plant"] == "PFB"]["ml"], 0.6)
    b_rural = [M.bvalue(allc[(allc["plant"] == p)]["ml"], 0.6)[0] for p in ("MOS", "WDR", "OTM")]
    ok(0.7 <= b_pfb <= 1.3, "Pfeilbrunn's located b-value above ML 0.6 is unremarkable (%.2f)" % b_pfb)
    # injection does not explain the day deficit: Pfeilbrunn injects more by day than by night
    h = d["hourly"]
    th = M.local_parts(M.epoch(h["time_utc"]))[2]
    pf = h["plant"].to_numpy() == "PFB"
    day = (th >= 8) & (th < 18)
    ok(h["flow_l_s"].to_numpy()[pf & day].mean() > h["flow_l_s"].to_numpy()[pf & M.quiet(th)].mean(),
       "Pfeilbrunn injects more by day than in the quiet hours")
    # the separation line: an ask device mishandled on the main path leaves the call unchanged (the reports, read
    # as labelled, in place of the telemetry for every volume)
    rep = A.report_frame(target)
    rep = rep[rep["function"] == "injection"]
    op = {pp["operator"]: pp["plant"] for pp in A.plant_register(target)["plants"]}
    rv = {p: {} for p in W.PLANTS}
    for x in rep.itertuples():
        if x.sheet == "2025" and x.month == "2025-12":
            pass
        y, m = int(x.month[:4]), int(x.month[5:7])
        k = (y, (m - 1) // 3 + 1)
        rv[op[x.operator]][k] = rv[op[x.operator]].get(k, 0) + x.volume_m3
    tot = {p: sum(v for k, v in rv[p].items() if k in W.record_quarters()) for p in W.PLANTS}
    rr = {}
    for p in W.PLANTS:
        e = c[c["plant"] == p]
        if p == "LZR":
            e = e[e["t"] >= d["switch"]]
        n = M.quiet(e["h"].to_numpy()).sum() / M.SHARE if p == "PFB" else len(e)
        frac = 1.0 if p != "LZR" else sum(v for (y, q), v in vol[p].items() if (y, q) >= (2025, 3)) / sum(vol[p].values())
        rr[p] = n / (tot[p] * frac)
    alt = {p: {q: rr[p] * rv[p].get((2026, q) if q < 4 else (2025, 4), 0) for q in (1, 2, 3, 4)} for p in W.PLANTS}
    ok(M.over(alt) == GOLD, "separation: the reports read as labelled for every volume leave the call: %s" % M.over(alt))
    return dict(d=d, c=c, g=g, r0=r0, r1=r1, r2=r2)


def asks(target, world):
    gold = A.golden(target)
    QS = A.QS
    # ask 1: the golden against the meters' truth by the month each reports; every cracker figure off
    inj = world["inj"]
    truth = {}
    for x in inj[inj["role"] == "injection"].itertuples():
        k = (int(x.year), (int(x.month) - 1) // 3 + 1)
        if k in QS:
            truth[(x.plant, k)] = truth.get((x.plant, k), 0) + int(x.volume_reported)
    lab = A.volumes(target, "as_labelled")
    noo = A.volumes(target, "no_october")
    reg = A.volumes(target, "register_only")
    for p in W.PLANTS:
        for k in QS:
            v = gold[p][k]["volume_raw"]
            ok(v == truth[(p, k)], "ask 1 golden equals the metered truth (%s %s)" % (p, k))
            ok(int(v) % 100 != 50, "ask 1 clear of the hundred's rounding edge (%s %s)" % (p, k))
            ok(A.round100(lab[p][k]) != gold[p][k]["volume_m3"], "ask 1 cracker (as labelled) misses %s %s" % (p, k))
            if k == (2026, 3):
                ok(A.round100(noo[p][k]) != gold[p][k]["volume_m3"], "ask 1 shift without October misses %s Q3" % p)
        ok(A.round100(reg["MOS"][(2026, 1)]) < 0.7 * gold["MOS"][(2026, 1)]["volume_m3"], "ask 1 hazard: MST-Th4a dropped")
    rep = A.report_frame(target)
    s26 = rep[rep["sheet"] == "2026"]
    ok(all(pd.Timestamp(r).strftime("%Y-%m") == m for r, m in zip(s26["received"], s26["month"])),
       "ask 1 device: 2026 sheet received in the labelled month")
    old = rep[rep["sheet"] != "2026"]
    ok(all(pd.Timestamp(r).strftime("%Y-%m") == "%d-%02d" % A._shift(m, 1) for r, m in zip(old["received"], old["month"])),
       "ask 1 device: 2023 to 2025 sheets received the month after the labelled month")
    dec25 = rep[(rep["sheet"] == "2025") & (rep["month"] == "2025-12")].set_index("well")["volume_m3"]
    jan26 = rep[(rep["sheet"] == "2026") & (rep["month"] == "2026-01")].set_index("well")["volume_m3"]
    ok(dec25.sort_index().equals(jan26.sort_index()), "ask 1 device: December 2025 filed twice with the same volumes")
    ok((rep["well"] == "MST-Th4a").sum() > 0 and (rep["well"] == "MST-GT4").sum() == 0, "ask 1 hazard: legacy well id")
    # the telemetry corroborates within a few parts per thousand where the months are labelled as reported
    d = M.load(target)
    _, vdf = M.volumes(d)
    ts = (pd.Timestamp(0, tz="UTC") + pd.to_timedelta(vdf["t"], unit="s")).dt.tz_convert(W.LOCAL)
    tm = vdf.assign(ym=ts.dt.strftime("%Y-%m").to_numpy()).groupby(["plant", "ym"])["v"].sum()
    op = {pp["operator"]: pp["plant"] for pp in A.plant_register(target)["plants"]}
    r23 = rep[(rep["sheet"] != "2026") & (rep["function"] == "injection")].assign(plant=lambda x: x["operator"].map(op))
    rm = r23.groupby(["plant", "month"])["volume_m3"].sum()
    dev = max(abs(rm[k] / tm[k] - 1) for k in rm.index if tm.get(k, 0) > 5e4)
    ok(dev < 0.012, "ask 1: telemetry and reports agree to %.4f where labels are months reported" % dev)
    # ask 2: the golden against the system's true events; the cracker paths off
    au = world["auto"]
    au = au[au["author"] == "scautoloc@lgrt1"].sort_values(["ev_key", "created"]).groupby("ev_key").tail(1)
    hm = M.holes(pd.read_csv(os.path.join(target, R.WELL_REG)))
    pl, dist = M.zone(au["lat"].round(4), au["lon"].round(4), hm)
    au = au.assign(plant=pl, dist=dist)
    ok(not ((au["dist"] > 4.95) & (au["dist"] < 5.05) & (au["ml_auto"] >= 1.0)).any(), "ask 2: no final origin within 50 m of 5 km")
    au = au[(au["plant"] != "") & (au["ml_auto"] >= 1.0)]
    y, q, _ = M.local_parts(au["t"].to_numpy())
    tcount = au.assign(y=y, q=q).groupby(["plant", "y", "q"]).size()
    rows = A.automatic_counts(target, "rows")
    allr = A.automatic_counts(target, "all_rows")
    allt = A.automatic_counts(target, "all_latest")
    for p in W.PLANTS:
        for k in QS:
            ok(gold[p][k]["automatic"] == int(tcount.get((p, k[0], k[1]), 0)), "ask 2 golden equals the true events (%s %s)" % (p, k))
            ok(rows[p][k] != gold[p][k]["automatic"], "ask 2 cracker (origins counted) misses %s %s" % (p, k))
            ok(allr[p][k] != gold[p][k]["automatic"], "ask 2 every origin counted misses %s %s" % (p, k))
    ok(sum(allt[p][k] != gold[p][k]["automatic"] for p in W.PLANTS for k in QS) >= 13,
       "ask 2: the latest origin of each event (the analysts') misses 13 or more figures")
    for wsec in (0.8, 1.5, 3.0, 5.0, 8.5):
        for dk in (3.0, 5.0, 8.0, 1000.0):
            alt = A.automatic_counts(target, "golden", wsec, dk)
            ok(all(alt[p][k] == gold[p][k]["automatic"] for p in W.PLANTS for k in QS),
               "ask 2: grouping at %.1f s and %.0f km gives the golden counts" % (wsec, dk))
    ok(set(A.origin_frame(target)["author"]) == {"scautoloc@lgrt1", "ga@lgrt1", "ae@lgrt1"}, "ask 2: origin authors")
    # ask 3: the golden against an independent count from the simulation's epochs and gaps
    st, gp = world["stations"], world["gaps"]
    hz = gp[gp["channel"] == "HHZ"]
    a0, z0 = int(W.YEAR_START_2026.timestamp()), int(W.RECORD_END.timestamp())
    reg_ = A.plant_register(target)
    gs = A.network_hours(target, "gaps_only")
    vu = A.network_hours(target, "visits_utc")
    vl = A.network_hours(target, "visits_local")
    anyc = A.network_hours(target, "any_channel")
    for pp in reg_["plants"]:
        per = {}
        for sta in pp["monitoring_network"]:
            live = [(int(e["start"].timestamp()), int(e["end"].timestamp()) if e["end"] else z0) for e in st if e["station"] == sta]
            hh = hz[hz["station"] == sta]
            per[sta] = A._minus(sorted(live), list(zip(hh["s"], hh["e"])))
        for k, (qa, qb) in A.qbounds().items():
            tr = A._under(per, qa, qb) / 3600
            g = gold[pp["plant"]][k]
            ok(abs(g["hours_raw"] - tr) < 1e-6, "ask 3 golden equals the independent count (%s %s)" % (pp["plant"], k))
            ok(abs(g["hours_raw"] * 10 - np.floor(g["hours_raw"] * 10) - 0.5) > 0.01, "ask 3 clear of the rounding edge by 4 s (%s %s)" % (pp["plant"], k))
            for name, alt in (("gaps only", gs), ("visits as UTC", vu), ("visits as local", vl), ("any channel", anyc)):
                ok(A.round1(alt[pp["plant"]][k]) != g["hours_under_four"], "ask 3 %s misses %s %s" % (name, pp["plant"], k))
    ok(len(reg_["plants"][2]["monitoring_network"]) == 4 and "LZR4" not in reg_["plants"][2]["monitoring_network"],
       "ask 3: Lanzenried's 2026 network is four stations without LZR4")
    return gold


def hygiene(target, golden_dir, meta_path):
    files = sorted(os.listdir(target))
    ok(len(files) >= 10, "input gate: 10 or more files (%d)" % len(files))
    ok(len({f.rsplit(".", 1)[1] for f in files}) >= 3, "input gate: 3 or more formats")
    big = max(len(pd.read_csv(os.path.join(target, f))) if f.endswith(".csv") else 0 for f in files)
    import pyarrow.parquet as pq
    big = max(big, pq.ParquetFile(os.path.join(target, R.HOURLY)).metadata.num_rows)
    ok(big >= 25000, "input gate: a file of 25,000 rows or more (%d)" % big)
    meta = json.load(open(meta_path))
    ok(len(meta["distractor_files"]) >= 2 and all(f in files for f in meta["distractor_files"]), "input gate: two distractors")
    ok(sorted(os.listdir(golden_dir)) == ["amber_outlook_2027.docx", "tls_annex_2026.xlsx"], "goldens are the two files asked")
    # no em or en dash anywhere written, no persona outside the four drawn
    texts = {}
    for f in files:
        p = os.path.join(target, f)
        if f.endswith((".csv", ".md", ".json", ".xml")):
            texts[f] = open(p, encoding="utf-8").read()
        elif f.endswith((".xlsx", ".docx")):
            with zipfile.ZipFile(p) as z:
                texts[f] = " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml"))
        elif f.endswith(".pdf"):
            from pypdf import PdfReader
            texts[f] = " ".join(pg.extract_text() for pg in PdfReader(p).pages)
    for f in os.listdir(golden_dir):
        with zipfile.ZipFile(os.path.join(golden_dir, f)) as z:
            texts["golden/" + f] = " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml"))
    for f, t in texts.items():
        ok("—" not in t and "–" not in t, "no em or en dash in %s" % f)
        ok("python-docx" not in t and "openpyxl" not in t.lower() and "ReportLab" not in t, "no library name in %s" % f)
    # dates past the extract date only where a file looks forward (plans, the 2027 revision dates)
    for f in (R.CAT, R.AUTO, R.GAPS, R.FEDERAL, R.GW, R.WELL_REG):
        dates = re.findall(r"\b(20\d\d-\d\d(?:-\d\d)?)", texts[f])
        ok(not dates or max(dates) <= W.EXTRACT_DATE.isoformat(), "no date after the extract in %s" % f)
    # the bulletin carries the catalogue's own counts
    d = M.load(target)
    t = M.table(M.counted(d))
    btxt = texts[os.path.basename("tls_bulletin_2026Q3.pdf")]
    for p in W.PLANTS:
        row = "%s %d %d %d %d" % (W.NAME[p], t[p][(2025, 4)], t[p][(2026, 1)], t[p][(2026, 2)], t[p][(2026, 3)])
        ok(row in " ".join(btxt.split()), "bulletin row for %s matches the catalogue" % p)
    # identity: one row per event and per origin
    ok(d["cat"]["event_id"].is_unique, "catalogue event ids unique")
    ok(A.origin_frame(target)["origin_id"].is_unique, "origin ids unique")


def figures(golden_dir, fig_path, gold_asks, main):
    f = json.load(open(fig_path))
    from docx import Document
    doc = Document(os.path.join(golden_dir, "amber_outlook_2027.docx"))
    text = " ".join(p.text for p in doc.paragraphs)
    cells = [[c.text for c in r.cells] for r in doc.tables[1].rows]
    for row in cells[1:]:
        name = row[0]
        for q, v in zip(("Q1", "Q2", "Q3", "Q4"), row[1:]):
            ok(int(v) == f["forecast_2027"][name][q], "memo table %s %s matches figures.json" % (name, q))
    for k in ("pfb_located_2026Q1", "pfb_located_2025Q4", "pfb_located_record", "pfb_located_quiet_hours"):
        ok(" %d " % f[k] in " %s " % text.replace(",", " ").replace(".", " "), "memo states %s" % k)
    for k in ("pfb_rate_per_100k_m3_quiet", "pfb_rate_per_100k_m3_all_hours", "pfb_rate_ratio"):
        ok(f[k] in text, "memo states %s" % k)
    import openpyxl
    ws = openpyxl.load_workbook(os.path.join(golden_dir, "tls_annex_2026.xlsx"), data_only=True)["Annex 2026"]
    seen = 0
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not r[0]:
            continue
        p = {W.NAME[x]: x for x in W.PLANTS}[r[0]]
        k = (2026, int(r[1][1]))
        g = gold_asks[p][k]
        ok(r[2] == g["volume_m3"] and r[3] == g["automatic"] and abs(r[4] - g["hours_under_four"]) < 1e-9,
           "annex row %s %s matches the asks" % (p, k))
        seen += 1
    ok(seen == 15, "annex has 15 rows")


def run(target, golden_dir, world=None, meta_path=None):
    world = world or S.simulate_all()
    N[0] = 0
    m = main_call(target, world)
    g = asks(target, world)
    hygiene(target, golden_dir, meta_path)
    figures(golden_dir, os.path.join(os.path.dirname(os.path.abspath(__file__)), "figures.json"), g, m)
    print("checks: %d assertions passed (forks SF1 to SF10, ladder, margins, asks, separation, gates)" % N[0])


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    task = os.path.dirname(here)
    run(os.path.join(task, "target"), os.path.join(task, "golden"), meta_path=os.path.join(task, "metadata.json"))
