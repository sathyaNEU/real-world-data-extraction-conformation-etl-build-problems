"""Independent verifier. Reads only target/ and golden/ (no generator state, no import from the generator) and
recomputes on its own code path: the spontaneous-report count, the between-person rungs (matched non-users,
observed against expected, new users of the established medicine), the comparison of each new user's first month
with the same person's days 61 to 180 for all three signals, the call, and every figure the note and the table
carry. Exits non-zero on any disagreement."""
import csv
import os
import re
import sys

import numpy as np
import openpyxl
import pandas as pd
from docx import Document

EPOCH = pd.Timestamp("2018-01-01")
END = pd.Timestamp("2026-06-30")
SIGNALS = {"tasorexant": ("fall", "zopiclone", "Fall"), "corinurad": ("aki", "allopurinol", "Acute kidney injury"),
           "velaglutide": ("pancreatitis", "orlistat", "Pancreatitis acute")}
FAIL = []


def check(cond, msg):
    if not cond:
        FAIL.append(msg)


def read_bundle(t):
    B = {}
    B["civ"] = pd.read_csv(os.path.join(t, "civil_registration_extract.csv"), dtype=str, keep_default_na=False)
    for c in ("doedsdato", "udvandringsdato", "indvandringsdato"):
        B["civ"][c] = pd.to_datetime(B["civ"][c].replace("", None))
    B["civ"] = B["civ"].set_index("pnr")
    rx = pd.read_parquet(os.path.join(t, "prescription_registry_2018-2026.parquet"))
    rx["eksd"] = pd.to_datetime(rx["eksd"])
    B["rx"] = rx
    B["atc"] = pd.read_csv(os.path.join(t, "atc_ddd_index_2026.csv"), sep=";", dtype=str, keep_default_na=False)
    c = pd.read_parquet(os.path.join(t, "patient_registry_contacts_2018-2026.parquet"))
    B["c"] = c.sort_values(["pnr", "starttid", "kontakt_id"]).reset_index(drop=True)
    B["dx"] = pd.read_parquet(os.path.join(t, "patient_registry_diagnoses_2018-2026.parquet"))
    B["m"] = pd.read_csv(os.path.join(t, "matched_comparators.csv"), dtype=str)
    B["m"]["indeksdato"] = pd.to_datetime(B["m"]["indeksdato"])
    B["pop"] = pd.read_csv(os.path.join(t, "population_and_admissions_2022-2025.csv"))
    wb = openpyxl.load_workbook(os.path.join(t, "pv_signal_screen_2026Q3.xlsx"), read_only=True)
    rows = list(wb["Line listing"].iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows) if r and r[0] == "Case number")
    B["ll"] = pd.DataFrame([r for r in rows[h + 1:] if r and r[0]], columns=list(rows[h]))
    return B


def resident_year(B, pnr, start):
    """Resident on every day of the 365 days up to and including the start."""
    civ = B["civ"].reindex(pnr)
    a = start - pd.Timedelta(days=365)
    I, U = civ["indvandringsdato"].to_numpy(), civ["udvandringsdato"].to_numpy()
    a, b = a.to_numpy(), start.to_numpy()
    ok = np.ones(len(pnr), bool)
    for i in range(len(pnr)):
        imm, emi = I[i], U[i]
        hi_ = not pd.isna(imm)
        hu = not pd.isna(emi)
        if hu and not hi_:
            ok[i] = b[i] < emi
        elif hi_ and not hu:
            ok[i] = a[i] >= imm
        elif hi_ and hu and imm > emi:
            ok[i] = (b[i] < emi) or (a[i] >= imm)
        elif hi_ and hu:
            ok[i] = (a[i] >= imm) and (b[i] < emi)
    return ok


def starts(B, substance):
    codes = B["atc"][B["atc"]["substans"] == substance]["atc"].tolist()
    rx = B["rx"][B["rx"]["atc"].isin(codes)].copy()
    rx["row"] = np.arange(len(rx))
    pos = rx[rx["apk"] > 0].sort_values("eksd")
    neg = rx[rx["apk"] < 0].sort_values("eksd")
    if len(neg):
        hit = pd.merge_asof(neg[["pnr", "vnr", "eksd"]], pos[["pnr", "vnr", "eksd", "row"]], on="eksd",
                            by=["pnr", "vnr"], direction="backward")
        pos = pos[~pos["row"].isin(set(hit["row"].dropna().astype(int)))]
    pos = pos.sort_values(["pnr", "eksd"])
    gap = pos.groupby("pnr")["eksd"].diff().dt.days
    st = pos[gap.isna() | (gap > 365)][["pnr", "eksd"]].rename(columns={"eksd": "start"})
    keep = resident_year(B, st["pnr"].to_numpy(), st["start"])
    return st[keep].reset_index(drop=True)


def admissions(B):
    c = B["c"]
    lpr2 = (c["kilde"] == "LPR2").to_numpy()
    phys = c["kontakttype"].fillna("ALCA00").eq("ALCA00").to_numpy()
    stay = np.where(lpr2, c["pattype"].isin(["0", "3"]).to_numpy(), phys)
    prev_end = c["sluttid"].shift(1)
    same = c["pnr"].eq(c["pnr"].shift(1)).to_numpy()
    gap_min = (c["starttid"] - prev_end).dt.total_seconds().to_numpy() / 60
    prev_lpr2 = np.r_[False, lpr2[:-1]]
    prev_stay = np.r_[False, stay[:-1]]
    l3 = same & ~lpr2 & ~prev_lpr2 & stay & prev_stay & (gap_min >= 0) & (gap_min <= 240)
    l2 = same & lpr2 & prev_lpr2 & stay & prev_stay & (c["starttid"].dt.normalize() == prev_end.dt.normalize()).to_numpy()
    c = c.assign(adm=np.cumsum(~(l3 | l2)), is3=~lpr2, in2=c["pattype"].eq("0"), er2=c["pattype"].eq("3"),
                 ac3=c["prioritet"].eq("ATA1") & c["kontakttype"].eq("ALCA00"), hasend=c["sluttid"].notna())
    a = c.groupby("adm", sort=True).agg(pnr=("pnr", "first"), start=("starttid", "min"), end=("sluttid", "max"),
                                        n=("pnr", "size"), nend=("hasend", "sum"), lpr3=("is3", "any"),
                                        inpat2=("in2", "any"), er2=("er2", "any"), ac3=("ac3", "any"))
    dur = (a["end"] - a["start"]).dt.total_seconds() / 3600
    a["inpatient"] = a["inpat2"] | (a["lpr3"] & ((dur >= 12) | (a["nend"] < a["n"])))
    a["emergency"] = a["er2"] | a["ac3"]
    return c[["kontakt_id", "adm"]], a


def events(B, kind):
    cid, A = B["adm"]
    dx = B["dx"].merge(cid, on="kontakt_id")
    code = dx["diagnosekode"]
    act = dx["diagnosetype"].eq("A")
    if kind == "fall":
        inj = code.str.match(r"^D(S\d\d|T(0\d|1[0-4]))") & act
        fallc = code.str.match(r"^DW(0\d|1\d)")
        ok = (dx.assign(i=inj).groupby("adm")["i"].any() & dx.assign(f=fallc).groupby("adm")["f"].any())
        sel = ok[ok].index
        sel = sel[(A.loc[sel, "inpatient"] | A.loc[sel, "emergency"]).to_numpy()]
    else:
        pre = "DN17" if kind == "aki" else "DK85"
        ok = dx.assign(h=code.str.startswith(pre) & act).groupby("adm")["h"].any()
        sel = ok[ok].index
        sel = sel[A.loc[sel, "inpatient"].to_numpy()]
    e = A.loc[sel, ["pnr", "start"]].copy()
    e["day"] = e["start"].dt.normalize()
    return e[["pnr", "day"]]


def obs_end(B, pnr, start):
    civ = B["civ"].reindex(pnr)
    end = pd.Series(END, index=range(len(pnr)))
    d = civ["doedsdato"].reset_index(drop=True)
    u = civ["udvandringsdato"].reset_index(drop=True)
    s = pd.Series(start.to_numpy())
    end = end.where(d.isna() | (d > end), d)
    end = end.where(u.isna() | (u < s) | (u > end), u)
    return end


def counts(ev, frame, lo_days, hi_days, end_col="end"):
    """Events of each row of frame (pnr, start, end) in [start + lo, min(start + hi, end)]; observed days."""
    f = frame.reset_index(drop=True).copy()
    f["i"] = np.arange(len(f))
    f["lo"] = f["start"] + pd.to_timedelta(lo_days, unit="D")
    f["hi"] = (f["start"] + pd.to_timedelta(hi_days, unit="D")).where(
        f["start"] + pd.to_timedelta(hi_days, unit="D") <= f[end_col], f[end_col])
    j = f[["i", "pnr", "lo", "hi"]].merge(ev, on="pnr")
    j = j[(j["day"] >= j["lo"]) & (j["day"] <= j["hi"])]
    n = j.groupby("i").size().reindex(range(len(f)), fill_value=0).to_numpy()
    days = ((f["hi"] - f["lo"]).dt.days + 1).clip(lower=0).to_numpy()
    return n, days


def cond_poisson(a, b, r, c):
    use = (a + b > 0) & (r > 0) & (c > 0)
    a, b, r, c = a[use].astype(float), b[use].astype(float), r[use].astype(float), c[use].astype(float)
    n = a + b

    def score(beta):
        return a.sum() - (n * r * np.exp(beta) / (r * np.exp(beta) + c)).sum()
    lo, hi = -6.0, 6.0
    for _ in range(200):
        mid = (lo + hi) / 2
        if score(mid) > 0:
            lo = mid
        else:
            hi = mid
    beta = (lo + hi) / 2
    p = r * np.exp(beta) / (r * np.exp(beta) + c)
    se = 1 / np.sqrt((n * p * (1 - p)).sum())
    return np.exp(beta), np.exp(beta - 1.959964 * se), np.exp(beta + 1.959964 * se)


def self_controlled(B, substance, subset=None):
    st = B["starts"][substance].copy()
    st = st.sort_values(["pnr", "start"]).reset_index(drop=True)
    nxt = st.groupby("pnr")["start"].shift(-1)
    e = obs_end(B, st["pnr"].to_numpy(), st["start"])
    e = e.where(nxt.isna() | (e < nxt - pd.Timedelta(days=1)), nxt - pd.Timedelta(days=1))
    st["end"] = e.to_numpy()
    if subset is not None:
        st = st[subset(st)].reset_index(drop=True)
    ev = B["events"][SIGNALS[substance][0]]
    a, r = counts(ev, st, 1, 30)
    b, c = counts(ev, st, 61, 180)
    return cond_poisson(a, b, r, c)


def spans(B, pnrs, frm=pd.Timestamp("2018-01-01")):
    """Observed spans of each person from frm: resident, alive, to 30 June 2026 (time abroad left out)."""
    civ = B["civ"].reindex(pnrs)
    out = []
    for p, d, u, i in zip(pnrs, civ["doedsdato"], civ["udvandringsdato"], civ["indvandringsdato"]):
        end = min(END, d) if not pd.isna(d) else END
        if not pd.isna(u) and not pd.isna(i) and i > u:      # left and came back
            out.append((p, frm, min(end, u - pd.Timedelta(days=1))))
            out.append((p, max(frm, i), end))
        else:
            lo = max(frm, i) if not pd.isna(i) else frm
            hi = min(end, u - pd.Timedelta(days=1)) if not pd.isna(u) else end
            out.append((p, lo, hi))
    sp = pd.DataFrame(out, columns=["pnr", "lo", "hi"])
    return sp[sp["hi"] >= sp["lo"]].reset_index(drop=True)


def all_time(B, substance, age=False):
    """Days 1 to 30 after every start against all other observed time from 1 January 2018: cases only, each person a
    stratum (age=True: each person and five-year age band a stratum, Mantel-Haenszel)."""
    st = B["starts"][substance]
    ev = B["events"][SIGNALS[substance][0]]
    sp = spans(B, sorted(set(st["pnr"]) & set(ev["pnr"])))
    hit = sp.reset_index().merge(ev, on="pnr")
    hit = hit[(hit["day"] >= hit["lo"]) & (hit["day"] <= hit["hi"])]
    sp = sp[sp["pnr"].isin(set(hit["pnr"]))]
    st = st[st["pnr"].isin(set(sp["pnr"]))]
    cut_at = {p: sorted(set(g["start"] + pd.Timedelta(days=1)) | set(g["start"] + pd.Timedelta(days=31)))
              for p, g in st.groupby("pnr")}
    years = list(pd.date_range("2018-01-01", "2027-01-01", freq="YS"))
    pieces = []
    for p, lo, hi in zip(sp["pnr"], sp["lo"], sp["hi"]):
        edges = sorted(set(cut_at.get(p, [])) | set(years))
        pts = [lo] + [e for e in edges if lo < e <= hi] + [hi + pd.Timedelta(days=1)]
        for a_, b_ in zip(pts[:-1], pts[1:]):
            pieces.append((p, a_, b_ - pd.Timedelta(days=1)))
    pc = pd.DataFrame(pieces, columns=["pnr", "lo", "hi"])
    risk = st.assign(rlo=st["start"] + pd.Timedelta(days=1), rhi=st["start"] + pd.Timedelta(days=30))
    j = pc.reset_index().merge(risk[["pnr", "rlo", "rhi"]], on="pnr")
    inrisk = set(j[(j["lo"] >= j["rlo"]) & (j["hi"] <= j["rhi"])]["index"])
    pc["risk"] = pc.index.isin(inrisk)
    pc["days"] = (pc["hi"] - pc["lo"]).dt.days + 1
    e = pc.reset_index().merge(ev, on="pnr")
    e = e[(e["day"] >= e["lo"]) & (e["day"] <= e["hi"])]
    pc["n"] = e.groupby("index").size().reindex(pc.index, fill_value=0).to_numpy()
    civ = B["civ"].reindex(pc["pnr"])
    pc["band"] = (pc["lo"].dt.year.to_numpy() - civ["foedselsaar"].astype(int).to_numpy()) // 5
    key = ["pnr", "band"] if age else ["pnr"]
    g = pc.groupby(key + ["risk"])[["n", "days"]].sum().unstack("risk", fill_value=0)
    a_, b_ = g[("n", True)].to_numpy(float), g[("n", False)].to_numpy(float)
    r_, c_ = g[("days", True)].to_numpy(float), g[("days", False)].to_numpy(float)
    if not age:
        return cond_poisson(a_, b_, r_, c_)[0]
    t = r_ + c_
    use = t > 0
    return float((a_ * c_ / np.where(use, t, 1))[use].sum() / (b_ * r_ / np.where(use, t, 1))[use].sum())


def first_month(B, kind, frame):
    f = frame.reset_index(drop=True).copy()
    f["end"] = obs_end(B, f["pnr"].to_numpy(), f["start"]).to_numpy()
    return counts(B["events"][kind], f, 1, 30)


def main(target, golden):
    B = read_bundle(target)
    B["adm"] = admissions(B)
    B["events"] = {k: events(B, k) for k in ("fall", "aki", "pancreatitis")}
    B["starts"] = {s: starts(B, s) for s in list(SIGNALS) + ["zopiclone", "allopurinol", "orlistat"]}
    lo22, last = pd.Timestamp("2022-01-01"), pd.Timestamp("2026-05-31")
    out = {}
    # rung 0: reports
    ll = B["ll"].sort_values(["Case number", "Version"]).drop_duplicates("Case number", keep="last")
    rep = {s: int((ll["Suspect"].str.lower().str.contains(s) &
                    ll["Reactions (MedDRA PT)"].apply(lambda x: SIGNALS[s][2] in [y.strip() for y in str(x).split(";")])).sum())
           for s in SIGNALS}
    check(max(rep, key=rep.get) == "tasorexant", "rung 0 winner %s" % rep)
    # the between-person rungs
    rungs = {}
    for s, (kind, comp, _) in SIGNALS.items():
        st = B["starts"][s]
        st = st[(st["start"] >= lo22) & (st["start"] <= last)]
        m = B["m"][B["m"]["substans"] == s]
        users = m[m["rolle"] == "bruger"].merge(st, left_on=["pnr", "indeksdato"], right_on=["pnr", "start"])
        mm = m[m["saet_id"].isin(set(users["saet_id"]))].rename(columns={"indeksdato": "start"})
        n, t = first_month(B, kind, mm[["pnr", "start"]])
        u = (mm["rolle"] == "bruger").to_numpy()
        matched = (n[u].sum() / t[u].sum()) / (n[~u].sum() / t[~u].sum())
        ac = B["starts"][comp]
        ac = ac[(ac["start"] >= lo22) & (ac["start"] <= last)]
        n1, t1 = first_month(B, kind, st)
        n0, t0 = first_month(B, kind, ac)
        active = (n1.sum() / t1.sum()) / (n0.sum() / t0.sum())
        # observed against expected at the national rates
        civ = B["civ"].reindex(st["pnr"])
        yr = st["start"].dt.year.to_numpy()
        age = yr - civ["foedselsaar"].astype(int).to_numpy()
        sex = civ["koen"].to_numpy()
        col = {"fall": "indlaeggelser_fald", "aki": "indlaeggelser_akut_nyreskade",
               "pancreatitis": "indlaeggelser_akut_pancreatitis"}[kind]
        p = B["pop"].assign(lo=B["pop"]["aldersgruppe"].str.extract(r"^(\d+)")[0].astype(int))
        lut = {(int(a), int(b), c_): v for a, b, c_, v in zip(p["aar"], p["lo"], p["koen"], p[col] / p["folketal_1_januar"])}
        bands = sorted(p["lo"].unique())
        exp = 0.0
        for i in range(len(st)):
            # a starter younger than the first band takes the first band's rate
            band = max([b_ for b_ in bands if b_ <= age[i]] or [bands[0]])
            exp += lut[(min(int(yr[i]), 2025), band, sex[i])] * t1[i] / 365.25
        rungs[s] = dict(matched=matched, active=active, oe=n1.sum() / exp)
    for name in ("matched", "oe", "active"):
        sc = {s: rungs[s][name] for s in SIGNALS}
        order = sorted(sc, key=lambda k: -sc[k])
        check(order[0] == "tasorexant" and sc[order[0]] / sc[order[1]] >= 1.3, "rung %s: %s" % (name, sc))
    check(rungs["corinurad"]["active"] <= 1.4, "corinurad against allopurinol starters %.3f" % rungs["corinurad"]["active"])
    out["rungs"] = rungs
    # the self-controlled comparison and the call
    sc = {s: self_controlled(B, s) for s in SIGNALS}
    supported = [s for s in SIGNALS if sc[s][1] > 1]
    call = max(supported, key=lambda s: sc[s][0])
    check(call == "corinurad", "the call %s %s" % (call, sc))
    rr, lo, hi = sc[call]
    # 2025 starters, admissions in days 1 to 30, and how many corinurad caused
    st = B["starts"][call]
    s25 = st[st["start"].dt.year == 2025]
    n25, _ = first_month(B, SIGNALS[call][0], s25)
    k = int(n25.sum())
    caused = int(np.floor(k * (rr - 1) / rr + 0.5))
    # NSAID-free starters
    ns = B["rx"][B["rx"]["atc"].str.startswith("M01A") & (B["rx"]["apk"] > 0)][["pnr", "eksd"]]

    def no_nsaid(frame):
        j = frame.reset_index().merge(ns, on="pnr")
        j = j[(j["eksd"] >= j["start"]) & (j["eksd"] <= j["start"] + pd.Timedelta(days=30))]
        return ~frame.index.isin(set(j["index"]))
    nf = self_controlled(B, call, subset=no_nsaid)
    # the year before the start
    def pre_rate(kind, frame):
        f = frame[(frame["start"] >= lo22) & (frame["start"] <= last)].reset_index(drop=True).copy()
        ok = resident_year(B, f["pnr"].to_numpy(), f["start"] - pd.Timedelta(days=1))
        f = f[ok].reset_index(drop=True)
        f["end"] = f["start"] - pd.Timedelta(days=1)
        n, _ = counts(B["events"][kind], f, -365, -1)
        return n.sum() / (len(f) * 365 / 365.25)
    mt = B["m"][(B["m"]["substans"] == "tasorexant") & (B["m"]["rolle"] == "sammenligning")].rename(
        columns={"indeksdato": "start"})[["pnr", "start"]]
    pre = dict(tas_matched=pre_rate("fall", B["starts"]["tasorexant"]) / pre_rate("fall", mt),
               tas_zop=pre_rate("fall", B["starts"]["tasorexant"]) / pre_rate("fall", B["starts"]["zopiclone"]),
               cor_all=pre_rate("aki", B["starts"]["corinurad"]) / pre_rate("aki", B["starts"]["allopurinol"]))
    # rung 2b: all observed time, the years before the start included, names velaglutide
    at = {s: all_time(B, s) for s in SIGNALS}
    at_age = {s: all_time(B, s, age=True) for s in SIGNALS}
    for name, sc in (("unadjusted", at), ("age bands within person", at_age)):
        order = sorted(sc, key=lambda k: -sc[k])
        check(order[0] == "velaglutide" and sc[order[0]] / sc[order[1]] >= 1.2, "all-time %s: %s" % (name, sc))
    # velaglutide is not started after a pancreatitis: its starters' year before, against orlistat starters'
    vel_pre = pre_rate("pancreatitis", B["starts"]["velaglutide"]) / pre_rate("pancreatitis", B["starts"]["orlistat"])
    vs = B["starts"]["velaglutide"]
    os_ = B["starts"]["orlistat"]

    def year_after(frame):
        f = frame[(frame["start"] >= lo22) & (frame["start"] <= last)].reset_index(drop=True).copy()
        f["end"] = obs_end(B, f["pnr"].to_numpy(), f["start"]).to_numpy()
        n, t = counts(B["events"]["pancreatitis"], f, 1, 365)
        return n.sum() / t.sum()
    vel_post = year_after(vs) / year_after(os_)
    check(vel_pre <= 0.25 and 0.75 <= vel_post <= 1.3, "velaglutide year before %.3f, year after %.3f" % (
        vel_pre, vel_post))
    # the goldens
    doc = Document(os.path.join(golden, "signal_refinement_note.docx"))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells)
    check("Open the formal assessment on corinurad (Uralexa) and acute kidney injury" in text, "the note's title")
    check("%.2f (%.2f to %.2f)" % (rr, lo, hi) in text, "ratio and interval %.2f (%.2f to %.2f) in the note" % (rr, lo, hi))
    check("ran at %.2f times the rate in the same people's days 61 to 180 (95%% CI %.2f to %.2f)" % (rr, lo, hi) in text,
          "the opening paragraph's ratio")
    check("%d admissions with acute kidney injury fell in that month, and corinurad caused %d of them" % (k, caused) in text,
          "2025 admissions %d and caused %d in the note" % (k, caused))
    for s in SIGNALS:
        check("%.2f (%.2f to %.2f)" % sc[s] in text, "%s ratio in the note" % s)
    check(re.search(r"at %.2f times the rate of the matched non-users" % pre["tas_matched"], text) is not None,
          "tasorexant against matched non-users %.2f" % pre["tas_matched"])
    check("at %.2f times the rate of people starting zopiclone" % pre["tas_zop"] in text,
          "tasorexant against zopiclone %.2f" % pre["tas_zop"])
    check("was %.2f times that of people starting allopurinol" % pre["cor_all"] in text,
          "corinurad against allopurinol %.2f" % pre["cor_all"])
    check("rate ratio is %.2f (95%% CI %.2f to %.2f)" % nf in text, "NSAID-free ratio %.2f" % nf[0])
    check("acute pancreatitis at %.2f times the rate of people starting orlistat" % vel_pre in text,
          "velaglutide year before %.2f in the note" % vel_pre)
    check("their rate is %.2f times orlistat starters'" % vel_post in text, "velaglutide year after %.2f" % vel_post)
    for s in SIGNALS:
        check(str(rep[s]) in text, "%s Q3 cases %d in the note" % (s, rep[s]))
    rows = list(csv.DictReader(open(os.path.join(golden, "new_users_by_year.csv"), encoding="utf-8")))
    check(len(rows) == 12, "twelve rows")
    for s in SIGNALS:
        st = B["starts"][s]
        for y in (2022, 2023, 2024, 2025):
            want = int(st[st["start"].dt.year == y]["pnr"].nunique())
            got = [int(r["new_users"]) for r in rows if r["substance"] == s and int(r["calendar_year"]) == y]
            check(got == [want], "%s %d: %s against %d" % (s, y, got, want))
    out.update(call=call, sccs=sc, ratio=(rr, lo, hi), events_2025=k, caused=caused, nsaid_free=nf, pre=pre, reports=rep,
               all_time=at, all_time_age=at_age, vel_pre=vel_pre, vel_post=vel_post)
    if FAIL:
        for f in FAIL:
            print("VERIFY FAIL:", f)
        raise SystemExit(1)
    print("verify OK: call %s, rate ratio %.2f (%.2f to %.2f), 2025 admissions %d, caused %d; tasorexant %.2f, "
          "velaglutide %.2f; NSAID-free %.2f; rungs %s; all observed time %s, within age bands %s; velaglutide year "
          "before %.2f and after %.2f x orlistat starters; new users table and note agree" % (
              call, rr, lo, hi, k, caused, sc["tasorexant"][0], sc["velaglutide"][0], nf[0],
              {s: {k_: round(v, 2) for k_, v in r.items()} for s, r in rungs.items()},
              {s: round(v, 2) for s, v in at.items()}, {s: round(v, 2) for s, v in at_age.items()}, vel_pre, vel_post))
    return out


if __name__ == "__main__":
    task = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    main(os.path.join(task, "target"), os.path.join(task, "golden"))
