"""Read the shipped pack back into frames (used by checks.py; verify.py has its own reader)."""
import datetime as dt
import os

import openpyxl
import pandas as pd

import files as FL
import world as W


def kestwick(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out, isins = {}, {}
    names = {W.FUNDS[f]["short"]: f for f in W.FUNDS}
    for ws in wb.worksheets:
        f = names[ws.title]
        line = ws["A2"].value
        parts = line.split()
        isins[f] = (parts[3], parts[7])
        rows = []
        for r in ws.iter_rows(min_row=6, values_only=True):
            if r[0] is None:
                continue
            rows.append((r[0].date() if isinstance(r[0], dt.datetime) else r[0], r[1], r[2], r[3] or "", r[4]))
        out[f] = pd.DataFrame(rows, columns=["Date", "Income units (GBP)", "Accumulation units (GBP)", "Ex-dividend",
                                             "Distribution per income unit (GBP)"])
    return out, isins


def budget(path):
    wb = openpyxl.load_workbook(path)
    rc = wb["Running costs"]
    running, capital, block = [], [], "run"
    for r in rc.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        if str(r[0]).startswith("Charged to capital"):
            block = "cap"
            continue
        if str(r[0]).startswith("Total"):
            continue
        (running if block == "run" else capital).append((r[0], r[1]))
    return running, capital


CSV_KEYS = ("val", "moves", "events", "ent", "static", "cal", "index", "reg", "inst", "pays")
DTYPES = {"lot_id": str, "lot_open_date": str, "sedol": str, "security_code": str, "new_security_code": str, "isin": str}


def read_frame(src):
    df = pd.read_csv(src, dtype=DTYPES, keep_default_na=False, na_values={"price": [""], "quantity": [""]})
    if "entitlement_gbp" in df.columns:
        df["entitlement_gbp"] = df["entitlement_gbp"].astype(float)
    return df


def pack(target):
    p = lambda k: os.path.join(target, FL.FILES[k])
    F = {}
    for k in CSV_KEYS:
        F[k] = read_frame(p(k))
    F["kest"], F["kest_isins"] = kestwick(p("kest"))
    F["running"], F["capital"] = budget(p("budget"))
    return F


def typed(F, keys=CSV_KEYS):
    """the generator's frames as the shipped files read back: each CSV frame written and read again exactly as
    files.write_csv and pack() do, Kestwick's history as its sheets carry it"""
    import io
    out = dict(F)
    for k in keys:
        buf = io.StringIO()
        F[k].to_csv(buf, index=False, lineterminator="\n")
        buf.seek(0)
        out[k] = read_frame(buf)
    out["kest_isins"] = {f: (W.FUNDS[f]["isin_inc"], W.FUNDS[f]["isin_acc"]) for f in W.FUNDS}
    return out
