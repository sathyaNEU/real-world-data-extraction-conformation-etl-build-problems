"""Independent verifier: reads only the shipped bundle (target/) and recomputes the call, the default path's call and
every graded ask figure with its own code, then compares with figures.json. Imports nothing from the generator.

    python3 task135/generator/verify.py [--target DIR]
"""
import argparse
import datetime as dt
import json
import os
import re
import sys

import numpy as np
import openpyxl
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
VAL = "2026-11-30"
INC_FROM = "2025-12-01"
STEPS = [20, 30, 40, 50, 60, 70, 80, 90, 100]
HOLD = ["MAND", "NA", "EU", "JP", "GILT", "CORP"]
SHEETS = {"North American Equity": "NA", "European Equity": "EU", "Japan Equity": "JP", "Gilt": "GILT",
          "Sterling Corporate Bond": "CORP"}
# investment policy 2.2
GROWTH = {"MAND": 0.40, "NA": 0.30, "EU": 0.15, "JP": 0.15}
DEFENSIVE = {"GILT": 0.55, "CORP": 0.45}
EUR = 1.1460            # finance procedure 7.5
FAILS = []


def check(cond, msg):
    print(("ok    " if cond else "FAIL  ") + msg)
    if not cond:
        FAILS.append(msg)


def weights(step):
    s = step / 100.0
    w = [GROWTH.get(h, 0) * s + DEFENSIVE.get(h, 0) * (1 - s) for h in HOLD]
    return np.array(w)


def r1000(x):
    return int(np.floor(x / 1000.0 + 0.5)) * 1000


def read(t):
    csv = lambda name: pd.read_csv(os.path.join(t, name), dtype=str, keep_default_na=False)
    X = dict(val=csv("brackenfold_valuation_lots_2024-11_2026-11.csv"), mov=csv("brackenfold_capital_movements_2024-12_2026-11.csv"),
             ev=csv("brackenfold_corporate_events_2024-12_2026-11.csv"),
             ent=csv("brackenfold_income_entitlements_by_lot_2025-11_2026-11.csv"),
             st=csv("brackenfold_static_data_2026-12-02.csv"), reg=csv("wrexford_grants_register_2026-11-30.csv"),
             ins=csv("wrexford_grant_instalments_2026-11-30.csv"), pay=csv("wrexford_grant_payments_2023-01_2026-11.csv"))
    for k in ("market_value_gbp", "accrued_income_gbp"):
        X["val"][k] = X["val"][k].astype(float)
    wb = openpyxl.load_workbook(os.path.join(t, "kestwick_charity_funds_prices_2024-11_2026-11.xlsx"), read_only=True)
    X["kw"], X["kw_isin"] = {}, {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        code = SHEETS[ws.title]
        X["kw_isin"].update({i: code for i in re.findall(r"GB00[A-Z0-9]{8}", rows[1][0])})
        hdr = rows[4]
        body = [r for r in rows[5:] if r[0] is not None]
        df = pd.DataFrame(body, columns=hdr)
        df["date"] = [x.date().isoformat() if hasattr(x, "date") else str(x) for x in df[hdr[0]]]
        df["price"] = df["Income units (GBP)"].astype(float)
        df["dist"] = df["Distribution per income unit (GBP)"].fillna(0).astype(float)
        X["kw"][code] = df.set_index("date")
    wb = openpyxl.load_workbook(os.path.join(t, "wrexford_budget_2027_draft_2026-11-24.xlsx"), read_only=True)
    lines = []
    for r in wb["Running costs"].iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        if str(r[0]).startswith("Total running costs"):
            break
        lines.append(float(r[1]))
    X["running"] = sum(lines)
    return X


def index_dates(t):
    """the benchmark index file's dates for each index, read under the header block above its column row"""
    path = os.path.join(t, "benchmark_index_levels_2024-11_2026-11.csv")
    with open(path, encoding="utf-8") as f:
        k = next(i for i, line in enumerate(f) if line.startswith("date,"))
    df = pd.read_csv(path, dtype=str, keep_default_na=False, skiprows=k)
    return {nm: set(s["date"]) for nm, s in df.groupby("index_name")}


def returns(X):
    v = X["val"]
    mand = v[v["asset_class"].isin(["Equity - UK", "Cash", "Rights"])]
    V = (mand.groupby("valuation_date")["market_value_gbp"].sum() + mand.groupby("valuation_date")["accrued_income_gbp"].sum())
    # the demerger: the new shares, credited a day after the parent went ex, valued on the ex-date at their first price
    sp = X["ev"][X["ev"]["action_type"] == "SOFF"].iloc[0]
    first = v[(v["security_code"] == sp["new_security_code"]) & (v["valuation_date"] == sp["effective_date"])]
    V.loc[sp["ex_date"]] += first["market_value_gbp"].sum()
    m = X["mov"][X["mov"]["movement_type"].isin(["Contribution", "Withdrawal"])].copy()
    m["a"] = m["amount_gbp"].astype(float)
    flows = m.groupby("date")["a"].sum().reindex(V.index).fillna(0.0)
    dates = list(V.index[1:])
    r = {"MAND": (V.values[1:] - flows.values[1:]) / V.values[:-1] - 1}
    for h in HOLD[1:]:
        k = X["kw"][h].reindex([V.index[0]] + dates)
        r[h] = (k["price"].values[1:] + k["dist"].values[1:]) / k["price"].values[:-1] - 1
    R = np.column_stack([r[h] for h in HOLD])
    return R, [dt.date.fromisoformat(d) for d in dates]


def vol_daily(R, step, lags=0, bartlett=False):
    p = R @ weights(step)
    y = p - p.mean()
    n = len(y)
    var = y @ y / (n - 1)
    for L in range(1, lags + 1):
        w = 1 - L / (lags + 1) if bartlett else 1.0
        var += 2 * w * (y[L:] @ y[:-L]) / (n - 1)
    return 100 * np.sqrt(252 * var)


def vol_weekly_fri(R, dates, step):
    """Friday-to-Friday returns: each holding's return over the week, combined at the step's policy weights (policy 4.2,
    "combined in every period"), weeks wholly inside the window, 52 a year"""
    wk = {}
    for x, d in zip(R, dates):
        fri = d + dt.timedelta(days=(4 - d.weekday()) % 7)
        wk.setdefault(fri, []).append(x)
    last = dates[-1]
    W = np.array([np.prod(1 + np.array(xs), axis=0) - 1 for fri, xs in sorted(wk.items()) if fri <= last])
    return 100 * np.sqrt(52 * np.var(W @ weights(step), ddof=1))


def vol_blocks(R, step, k=5):
    """consecutive five-day periods from the window's first day, each holding compounded, combined at the policy weights"""
    n = len(R) // k * k
    B = (1 + R[:n]).reshape(-1, k, R.shape[1]).prod(axis=1) - 1
    return 100 * np.sqrt(52 * np.var(B @ weights(step), ddof=1))


def call(vols):
    ok = [s for s in STEPS if vols[s] <= 10.0]
    return max(ok) if ok else None


def asks(X):
    v = X["val"]
    last = v[v["valuation_date"] == VAL]
    mand_mv = last[last["asset_class"].isin(["Equity - UK", "Cash", "Rights"])]["market_value_gbp"].sum()
    isin = dict(zip(X["st"]["security_code"], X["st"]["isin"]))
    fund_of = {c: X["kw_isin"][i] for c, i in isin.items() if i in X["kw_isin"]}
    fund_mv = {h: 0.0 for h in HOLD[1:]}
    for c, mv in zip(last["security_code"], last["market_value_gbp"]):
        if c in fund_of:
            fund_mv[fund_of[c]] += mv
    total = mand_mv + sum(fund_mv.values())
    e = X["ent"]
    e = e[(e["ex_date"] >= INC_FROM) & (e["ex_date"] <= VAL) & (~e["security_code"].isin(list(fund_of)))]
    one = e.drop_duplicates(["sub_account", "security_code", "event_ref"])
    y = {"MAND": one["entitlement_gbp"].astype(float).sum() / mand_mv}
    for h in HOLD[1:]:
        k = X["kw"][h]
        sel = k[(k.index >= INC_FROM) & (k.index <= VAL)]
        y[h] = sel["dist"].sum() / k.loc[VAL, "price"]
    inc = {s: total * sum(w * y[h] for w, h in zip(weights(s), HOLD)) for s in STEPS}
    # grant payments: open instalments of grants not withdrawn, at the date the trust expects to pay them
    reg = X["reg"].set_index("grant_ref")
    ins = X["ins"].copy()
    ins["due"] = [dt.datetime.strptime(x, "%d/%m/%Y").date() for x in ins["due_date"]]
    ins["kind"] = ["core" if reg.loc[g, "grant_type"] == "Core" else ("first" if n == "1" else "later")
                   for g, n in zip(ins["grant_ref"], ins["instalment_no"])]
    pay = X["pay"].merge(ins, on=["grant_ref", "instalment_no"])
    lag = pay[pay["kind"] == "later"].apply(lambda r: (dt.date.fromisoformat(r["payment_date"]) - r["due"]).days, axis=1)
    L = int(np.median(lag))
    G = 0.0
    for r in ins.itertuples(index=False):
        if r.status != "Scheduled" or reg.loc[r.grant_ref, "status"] == "Withdrawn":
            continue
        when = r.due + dt.timedelta(days=L) if r.kind == "later" else r.due
        if when.year == 2027:
            G += float(r.amount) / (EUR if reg.loc[r.grant_ref, "currency"] == "EUR" else 1.0)
    rel = {s: G + X["running"] - inc[s] for s in STEPS}
    title = max(s for s in STEPS if inc[s] >= G)
    expendable = last[last["sub_account"] == "WRX-EX02"]["market_value_gbp"].sum()
    return dict(total=total, yields=y, inc=inc, G=G, running=X["running"], rel=rel, title=title, lag=L,
                expendable=expendable)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--target", default=os.path.join(TASK, "target"))
    a = ap.parse_args()
    F = json.load(open(os.path.join(HERE, "figures.json")))
    X = read(a.target)
    R, dates = returns(X)
    check(R.shape == (505, 6) and dates[-1] == dt.date(2026, 11, 30), "505 daily returns for six holdings to 30 November 2026")
    r0 = {s: vol_daily(R, s) for s in STEPS}
    fams = {"lead-lag 1": {s: vol_daily(R, s, 1) for s in STEPS},
            "newey-west 4": {s: vol_daily(R, s, 4, True) for s in STEPS},
            "calendar weeks to Fri": {s: vol_weekly_fri(R, dates, s) for s in STEPS},
            "five-day blocks, anchor 0": {s: vol_blocks(R, s, 5) for s in STEPS}}
    check(call(r0) == 70 == F["call_default"], "the default (daily covariance, 252 days) calls %s (figures.json %s)" % (call(r0), F["call_default"]))
    for s in STEPS:
        check(abs(r0[s] - F["ladder"]["R0 daily"][str(s)]) < 1e-3, "R0 at %d: %.4f (figures.json %.4f)" % (s, r0[s], F["ladder"]["R0 daily"][str(s)]))
    for n, v in fams.items():
        check(call(v) == 60 == F["call"], "%s calls %s" % (n, call(v)))
        for s in (60, 70):
            check(abs(v[s] - F["ladder"]["G " + n][str(s)]) < 1e-3, "%s at %d: %.4f (figures.json %.4f)" % (n, s, v[s], F["ladder"]["G " + n][str(s)]))
    # an inner join of every price file on date (the mandate's valuation dates, Kestwick's sheets, the index file), each
    # return carried across any date the join drops, then the default path
    every = sorted(set(X["val"]["valuation_date"]))
    joined = set(every)
    for h in HOLD[1:]:
        joined &= set(X["kw"][h].index)
    ixd = index_dates(a.target)
    for s_ in ixd.values():
        joined &= s_
    rows, carry = [], np.ones(R.shape[1])
    for x, d in zip(R, dates):
        carry = carry * (1 + x)
        if d.isoformat() in joined:
            rows.append(carry - 1)
            carry = np.ones(R.shape[1])
    Rj = np.array(rows)
    rj = {s: vol_daily(Rj, s) for s in STEPS}
    check(len(ixd) == 6 and len(joined) == len(every) == 506 and len(Rj) == 505 and call(rj) == 70 == F["index_join"]["default"]
          and abs(rj[70] - F["index_join"]["r0_70"]) < 0.005,
          "inner join with the index file: %d of %d dates kept, %d returns, the default calls %s (%.2f at 70)"
          % (len(joined), len(every), len(Rj), call(rj), rj[70]))
    A = asks(X)
    check(abs(A["total"] - F["portfolio_value"]) < 0.01, "portfolio value %.2f" % A["total"])
    check(abs(A["running"] - F["running_costs"]) < 0.01, "running costs %.0f" % A["running"])
    check(abs(A["G"] - F["grant_payments_2027"]) < 0.01 and r1000(A["G"]) == F["grant_payments_2027_stated"],
          "2027 grant payments %.2f, stated %d (ledger median lag %d days)" % (A["G"], r1000(A["G"]), A["lag"]))
    for s in STEPS:
        check(abs(A["inc"][s] - F["income"][str(s)]) < 0.01 and r1000(A["inc"][s]) == F["income_stated"][str(s)],
              "income at %d: %.2f, stated %d" % (s, A["inc"][s], r1000(A["inc"][s])))
        check(abs(A["rel"][s] - F["release"][str(s)]) < 0.01 and r1000(A["rel"][s]) == F["release_stated"][str(s)],
              "release at %d: %.2f, stated %d" % (s, A["rel"][s], r1000(A["rel"][s])))
    check(A["title"] == F["title_step"], "chart title step %d" % A["title"])
    check(all(0 < A["rel"][s] < A["expendable"] for s in STEPS),
          "every release under the expendable endowment's %.2f at 30 November (largest %.2f): a cap at its value "
          "never binds (policy 6.1, 6.2)" % (A["expendable"], max(A["rel"].values())))
    print("verify: %d failed" % len(FAILS))
    return 0 if not FAILS else 1


if __name__ == "__main__":
    sys.exit(main())
