"""The two deliverables, written from generator/figures.json: growth_share_2027.xlsx and income_cost_by_step_2027.svg."""
import datetime as dt
import json
import os
import re
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import officex as OX

TASK = os.path.dirname(HERE)
STEPS = [20, 30, 40, 50, 60, 70, 80, 90, 100]
BLUE, ORANGE, VIOLET = "#2a78d6", "#eb6834", "#6250d6"
SURFACE, INK, INK2, GRID = "#fcfcfb", "#0b0b0b", "#52514e", "#e6e5e1"


def r1000(x):
    import math
    return int(math.floor(x / 1000.0 + 0.5)) * 1000


def workbook(fig, path):
    G = fig["grant_payments_2027_stated"]
    RC = fig["running_costs_stated"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2027 growth share"
    bold = Font(bold=True)
    ws["A1"] = "The Wrexford Trust: growth share and budget lines for 2027"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Annual investment review for the trustees' meeting on 10 December 2026. Valuation of 30 November 2026."
    ws["A4"] = "Largest growth share inside the 10% volatility cap"
    ws["A4"].font = bold
    ws["C4"] = fig["call"] / 100.0
    ws["C4"].number_format = "0%"
    ws["C4"].font = Font(bold=True, size=12)
    ws["A5"] = "2027 grant payments (GBP)"
    ws["C5"] = G
    ws["A6"] = "2027 running costs (GBP)"
    ws["C6"] = RC
    ws["A7"] = "Grant payments and running costs (GBP)"
    ws["C7"] = "=C5+C6"
    for c in ("C5", "C6", "C7"):
        ws[c].number_format = "#,##0"
    hdr = ["Growth share", "Expected volatility, daily returns with one-day lead and lag (%)",
           "Expected volatility, daily returns, lagged covariances to five days (%)", "Inside the cap",
           "2027 investment income (GBP)", "Capital release (GBP)"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=9, column=j, value=h)
        c.font = bold
        c.fill = PatternFill("solid", fgColor="DCE6F1")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ll = fig["ladder"]["G lead-lag 1"]
    wk = fig["ladder"]["G newey-west 5"]
    for k, s in enumerate(STEPS):
        r = 10 + k
        inside = ll[str(s)] <= 10.0 and wk[str(s)] <= 10.0
        ws.cell(row=r, column=1, value=s / 100.0).number_format = "0%"
        ws.cell(row=r, column=2, value=round(ll[str(s)], 1)).number_format = "0.0"
        ws.cell(row=r, column=3, value=round(wk[str(s)], 1)).number_format = "0.0"
        ws.cell(row=r, column=4, value="Yes" if inside else "No").alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=5, value=fig["income_stated"][str(s)]).number_format = "#,##0"
        ws.cell(row=r, column=6, value="=$C$7-E%d" % r).number_format = "#,##0"
        if s == fig["call"]:
            for j in range(1, 7):
                ws.cell(row=r, column=j).font = bold
    # the release formula reproduces the stated figure at every step (asserted: the parts round the same way)
    for s in STEPS:
        assert G + RC - fig["income_stated"][str(s)] == fig["release_stated"][str(s)], s
    ws["A20"] = "Volatility before rounding is compared with the cap (investment policy 4.4). Pounds to the nearest thousand; notes on the Notes sheet."
    ws["A20"].font = Font(italic=True, color="595959")
    for col, w in zip("ABCDEF", [46, 17, 17, 9, 17, 15]):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[9].height = 62
    ws.freeze_panes = "A10"
    nt = wb.create_sheet("Notes")
    nt["A1"] = "Notes"
    nt["A1"].font = bold
    fw = fig["forcing"]
    notes = [
        "Expected volatility is the standard deviation of the portfolio's twelve-month total return (investment policy "
        "4.1 to 4.4), estimated from the six holdings' total returns over the two years to the 30 November 2026 "
        "valuation, combined at each growth share's policy weights.",
        "The Kestwick funds are dealt at their 12 noon valuation point from each market's latest prices, so the North "
        "American fund's price on a day carries the previous evening's New York close; the Fernside shares are priced at "
        "the London close. The mandate's daily return correlates %.2f with the North American fund's on the same day "
        "and %.2f with the fund's next-day return, so the twelve-month variance takes in the co-movement that lands a day "
        "apart: the daily covariances with the one-day lead and lag added, and, as a check, the daily covariances "
        "lagged out to five days." % (fw["uk_na_same_day"], fw["uk_na_next_day"]),
        "At 70 per cent growth both estimates are over the cap (%.1f and %.1f per cent). The daily covariance on its own, "
        "scaled by the square root of 252, gives %.1f per cent at 70 per cent, because it misses the co-movement that "
        "lands a day apart." % (ll["70"], wk["70"], fig["default_r0"]["70"]),
        "Investment income (finance procedure 7.2 and 7.3): the mandate's dividends with ex-dividend dates from 1 December "
        "2025 to 30 November 2026, one amount per designation, security and event (the custodian's income report "
        "repeats each holding's entitlement on every tax lot line), over the mandate's market value at 30 November with "
        "its cash and without accrued income; each Kestwick fund's distributions per income unit over the same twelve "
        "months, over its 30 November price. Each yield is applied to the share the step's policy weights give the "
        "holding of the portfolio's market value, GBP {:,.0f}.".format(fig["portfolio_value"]),
        "Grant payments (procedure 7.4 and 7.5): the open instalments of grants awarded by the budget date, leaving out "
        "grants the register shows as withdrawn. Core grants and first instalments are paid on their due date; later "
        "project instalments are paid once the progress report is accepted, %d to %d days after the due date in every "
        "payment since January 2023, so those due from late November are paid in the following year. Euro grants at "
        "1.1460 euro to the pound." % (fig["lags"]["later_min"], fig["lags"]["later_max"]),
        "Running costs (procedure 7.6): the draft budget's running costs total, GBP {:,.0f}. The Fernside and Brackenfold "
        "fees are charged to capital (investment policy 5.1) and are not included.".format(fig["running_costs"]),
        "Capital release: grant payments plus running costs less investment income, the release from the expendable "
        "endowment under investment policy 6.1.",
    ]
    for k, t in enumerate(notes, 1):
        c = nt.cell(row=k + 1, column=1, value="%d. %s" % (k, t))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        nt.row_dimensions[k + 1].height = 15 * (len(t) // 120 + 1)
    nt.column_dimensions["A"].width = 130
    OX.xlsx_from_openpyxl(wb, path, "Harriet Clarke", "Harriet Clarke", dt.datetime(2026, 12, 3, 9, 12),
                          dt.datetime(2026, 12, 4, 16, 48), company="The Wrexford Trust", seed=1204)


def chart(fig, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    matplotlib.rcParams["svg.hashsalt"] = "wrexford-2027"
    matplotlib.rcParams["font.family"] = "DejaVu Sans"
    inc = [fig["income_stated"][str(s)] / 1000.0 for s in STEPS]
    g = fig["grant_payments_2027_stated"] / 1000.0
    gr = fig["grants_plus_running_stated"] / 1000.0
    t = fig["title_step"]
    fig_, ax = plt.subplots(figsize=(8.8, 5.0))
    fig_.patch.set_facecolor(SURFACE)
    ax.set_facecolor(SURFACE)
    x = list(range(len(STEPS)))
    ax.bar(x, inc, width=0.56, color=BLUE, label="2027 investment income", zorder=2)
    ax.axhline(g, color=ORANGE, lw=2, zorder=3, solid_capstyle="round", label="2027 grant payments")
    ax.axhline(gr, color=VIOLET, lw=2, zorder=3, solid_capstyle="round", label="Grant payments and running costs")
    lo = 0
    hi = max(inc + [gr]) * 1.08
    ax.set_ylim(lo, hi)
    ax.set_xlim(-0.6, len(STEPS) - 0.4)
    for k, v in enumerate(inc):
        if STEPS[k] in (20, t, 100):
            ax.text(k, v + (hi - lo) * 0.012, format(int(round(v * 1000)), ","), ha="center", va="bottom", fontsize=8,
                    color=INK)
    ax.text(len(STEPS) - 0.45, g + (hi - lo) * 0.012, "GBP " + format(int(round(g * 1000)), ","), ha="right",
            va="bottom", fontsize=8.5, color=INK)
    ax.text(len(STEPS) - 0.45, gr + (hi - lo) * 0.012, "GBP " + format(int(round(gr * 1000)), ","), ha="right",
            va="bottom", fontsize=8.5, color=INK)
    ax.set_xticks(x)
    ax.set_xticklabels(["%d%%" % s for s in STEPS])
    ax.set_xlabel("Growth share", color=INK2, fontsize=9.5)
    ax.set_ylabel("GBP thousands", color=INK2, fontsize=9.5)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, p: format(int(v), ",")))
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color("#c9c8c4")
    ax.tick_params(colors=INK2, labelsize=9)
    ax.yaxis.grid(True, color=GRID, lw=0.8, zorder=0)
    ax.set_axisbelow(True)
    h, lb = ax.get_legend_handles_labels()
    order = [lb.index("2027 investment income"), lb.index("2027 grant payments"), lb.index("Grant payments and running costs")]
    ax.legend([h[i] for i in order], [lb[i] for i in order], loc="upper center", bbox_to_anchor=(0.5, -0.13), ncol=3,
              frameon=False, fontsize=8.5, labelcolor=INK)
    ax.set_title("2027 income covers the grant payments up to %d%% growth" % t, loc="left", fontsize=12.5, color=INK,
                 pad=14)
    fig_.text(0.012, 0.012, "Investment income by finance procedure 7.2 at each growth share, valuation of 30 November 2026.\n"
                            "Grant payments and running costs for the 2027 budget. The Wrexford Trust, December 2026.",
              fontsize=7.2, color=INK2, linespacing=1.4)
    fig_.tight_layout(rect=(0, 0.05, 1, 1))
    fig_.savefig(path, format="svg", metadata={"Creator": None, "Date": None, "Format": None, "Type": None})
    plt.close(fig_)
    s = open(path, encoding="utf-8").read()
    s = re.sub(r"<metadata>.*?</metadata>\s*", "", s, flags=re.S)
    s = re.sub(r"<!-- Created with [^>]*-->\s*", "", s)
    s = s.replace('id="matplotlib.axis_1"', 'id="xaxis"').replace('id="matplotlib.axis_2"', 'id="yaxis"')
    open(path, "w", encoding="utf-8").write(s)


def build(out=None):
    fig = json.load(open(os.path.join(HERE, "figures.json")))
    out = out or os.path.join(TASK, "golden")
    os.makedirs(out, exist_ok=True)
    workbook(fig, os.path.join(out, "growth_share_2027.xlsx"))
    chart(fig, os.path.join(out, "income_cost_by_step_2027.svg"))


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else None)
