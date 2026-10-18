"""The program's exports as its systems write them: the laboratories' results loaded to the state database, the
samplers' collection records, the source facility inventory, the 2025 monthly operating reports, the meter
register, the system register and its wholesale connections, the laboratory list, and two files kept beside them
(the 2023 national survey and the private well program's results)."""
from __future__ import annotations

import csv
import datetime as dt
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import common as K
import world as W

LABS = [("SEHL", "State Environmental Health Laboratory", "Bredesen", "EPA 533", "9901"),
        ("TAI", "Tolvanen Analytical, Inc.", "Nyquist", "EPA 533", "4417"),
        ("KET", "Kestrel Environmental Testing", "Port Ellefsen", "EPA 537.1", "4532")]
M5371 = ["PFBS", "PFHxA", "PFHpA", "PFOA", "PFNA", "PFDA", "PFUnA", "PFDoA", "PFTrDA", "PFTA", "PFHxS", "PFOS",
         "NMeFOSAA", "NEtFOSAA", "HFPO-DA", "ADONA", "9Cl-PF3ONS", "11Cl-PF3OUdS"]
EXTRA_CAS = {"PFTrDA": "72629-94-8", "PFTA": "376-06-7", "NMeFOSAA": "2355-31-9", "NEtFOSAA": "2991-50-6"}
CAS = dict((a, c) for a, c, _ in K.ANALYTES)
CAS.update(EXTRA_CAS)
RL5371 = {a: (2.0 if a in ("HFPO-DA", "ADONA", "NMeFOSAA", "NEtFOSAA", "9Cl-PF3ONS", "11Cl-PF3OUdS") else 1.0)
          for a in M5371}


def _w(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)
    return len(rows)


def fmt(v):
    """A laboratory's three significant figures, as the database stores them."""
    if v is None:
        return ""
    s = "%.3g" % v
    return s if "e" not in s else "%.0f" % v


# ---------------------------------------------------------------------------------------------- who sampled, where
def assign(S):
    """Each system's laboratory and sampler, and the way that sampler writes the wells in service."""
    for s in S:
        if not s["eps"]:
            continue
        r = K.rng("assign", s["pws"])
        u = r.random()
        s["lab"] = "SEHL" if u < 0.6 else ("TAI" if u < 0.85 else "KET")
        s["sampler"] = "Program sampler %02d" % int(r.integers(1, 10)) if r.random() < 0.7 else \
            "Operator %d" % int(r.integers(10200, 19900))
        s["style"] = int(r.integers(0, 6))
    return S


def sis_text(r, style, ep, on, plant=False):
    """The collection record's sources in service, in the sampler's own hand."""
    if plant:
        return "Plant"
    nos = [str(n) for n in on]
    active = [w["no"] for w in ep["wells"]]
    if len(on) > 1 and sorted(on) == sorted(active) and r.random() < 0.35:
        return ["All wells", "all wells on", "All", "ALL WELLS", "all 3 wells" if len(on) == 3 else "both wells",
                "All wells running"][style]
    if len(on) == 1:
        return ["Well %s only" % nos[0], "Well %s" % nos[0], "W%s" % nos[0], "#%s" % nos[0], nos[0],
                "Well #%s" % nos[0]][style]
    if style == 0:
        return ", ".join("Well %s" % n for n in nos)
    if style == 1:
        return "Wells " + ", ".join(nos[:-1]) + " & " + nos[-1]
    if style == 2:
        return ", ".join("W%s" % n for n in nos)
    if style == 3:
        return " and ".join("#%s" % n for n in nos)
    if style == 4:
        return ",".join(nos)
    return " + ".join("Well %s" % n for n in nos)


def unique_flow(ep, on):
    """True when no other set of the entry point's wells has the same rated capacity in total (a blank record whose
    wells the entry-point flow still names)."""
    import itertools
    rate = {w["no"]: w["gpm"] for w in ep["wells"]}
    tot = sum(rate[n] for n in on)
    nos = sorted(rate)
    for k in range(1, len(nos) + 1):
        for combo in itertools.combinations(nos, k):
            if sorted(combo) != sorted(on) and abs(sum(rate[n] for n in combo) - tot) <= 0.06 * tot:
                return False
    return True


def number_samples(X):
    """The state sample numbers, given in the order the samples reached the database (by collection)."""
    X.sort(key=lambda x: (x["when"], x["pws"], x["ep"], ["RT", "CO", "FD", "FRB"].index(x["kind"])))
    seq = {}
    for x in X:
        y = x["when"].year % 100
        seq[y] = seq.get(y, 0) + 1 + int(K.rng("gap", y, seq.get(y, 0)).integers(0, 3) == 0)
        x["no"] = "PF%02d-%05d" % (y, seq[y])
    return X


def revisions(X):
    """The few reports a laboratory reissued: one result corrected, the whole report loaded again beside the
    first with the next revision number."""
    for x in X:
        x["revs"] = [x["res"]]
        if x["kind"] not in ("RT", "CO"):
            continue
        r = K.rng("rev", x["no"])
        if r.random() < 0.013:
            a = ["PFOA", "PFOS", "PFHxS", "PFBS", "PFHxA"][int(r.integers(0, 5))]
            new = dict(x["res"])
            v = x["true"][a] * r.uniform(0.9, 1.1)
            new[a] = float("%.3g" % v) if v >= K.RLS[a] else None
            x["revs"].append(new)
            x["rev_days"] = int(r.integers(9, 40))
    return X


def write_results(path, S, X):
    by = {s["pws"]: s for s in S}
    rows = []
    for x in X:
        s = by[x["pws"]]
        lab = s["lab"]
        method = "EPA 537.1" if lab == "KET" else "EPA 533"
        alist = M5371 if lab == "KET" else K.AN
        r = K.rng("labid", x["no"])
        lab_id = {"SEHL": "SE%s%05d" % (x["when"].strftime("%y"), int(r.integers(0, 99999))),
                  "TAI": "R%s-%04d-%02d" % (x["when"].strftime("%y%m"), int(r.integers(1, 9999)), int(r.integers(1, 12))),
                  "KET": "L%07d" % int(r.integers(1000000, 9999999))}[lab]
        analyzed = (x["when"] + dt.timedelta(days=int(r.integers(4, 19)))).date()
        for rev, res in enumerate(x["revs"], start=1):
            reported = analyzed + dt.timedelta(days=int(K.rng("rep", x["no"]).integers(2, 9)) +
                                               (x.get("rev_days", 0) if rev > 1 else 0))
            for a in alist:
                v = res.get(a) if a in res else None
                rl = RL5371[a] if lab == "KET" else K.RLS[a]
                if v is not None and v < rl:
                    v = None
                rows.append([x["no"], lab, lab_id, x["pws"], s["full"], x["ep"], x["kind"],
                             x["when"].date().isoformat(), x["when"].strftime("%H:%M"), method, a, CAS[a],
                             "<" if v is None else "=", fmt(v), "%.1f" % rl, "ng/L", analyzed.isoformat(),
                             reported.isoformat(), rev])
    return _w(path, ["state_sample_no", "lab_code", "lab_sample_id", "pws_id", "system_name", "sample_point",
                     "sample_type", "collected_date", "collected_time", "method", "analyte", "cas_rn", "result_sign",
                     "result", "reporting_limit", "units", "analyzed_date", "report_date", "report_revision"], rows)


FIELD_NOTES = ["tap flushed 5 min", "rain overnight", "operator on site", "new sample tap installed", "cooler 4 C at pickup",
               "residual by kit, analyzer offline", "road work at pumphouse", "flow read from SCADA screen",
               "pumphouse locked, operator let us in", "tap flushed 10 min, slight color cleared", "operator absent",
               "sampled before hydrant flushing", "bottle 2 cracked, refilled", "ice on lot, late arrival"]


def write_collection(path, S, X):
    by = {s["pws"]: s for s in S}
    rows = []
    visit = {}
    for x in X:
        s = by[x["pws"]]
        ep = next(e for e in s["eps"] if e["code"] == x["ep"])
        key = (x["pws"], x["ep"], x["when"].date(), x["kind"] in ("FD", "FRB"))
        r = K.rng("record", x["no"])
        if x["kind"] == "FD":
            text, flow, cl = visit[(x["pws"], x["ep"], x["when"].date())]
        elif x["kind"] == "FRB":
            text, flow, cl = "", "", ""
        else:
            plant = any(w.get("plant") for w in ep["wells"])
            text = sis_text(r, s["style"], ep, x["on"], plant)
            flow = int(5 * round(sum(x["flows"].values()) * (1 + r.normal(0, 0.012)) / 5))
            cl = "%.2f" % r.uniform(0.18, 0.95) if r.random() < 0.85 else ""
            if r.random() < 0.016 and not plant and unique_flow(ep, x["on"]):
                text = ""
            visit[(x["pws"], x["ep"], x["when"].date())] = (text, flow, cl)
        remark = ""
        if x["kind"] == "FRB":
            remark = "FRB opened at tap"
        elif x["kind"] == "FD":
            remark = "duplicate"
        elif not text:
            remark = "sources not recorded"
        elif x["kind"] == "CO":
            remark = "confirmation"
        note = FIELD_NOTES[int(r.integers(0, len(FIELD_NOTES)))] if r.random() < 0.13 else ""
        if note and x["kind"] in ("RT", "CO"):
            remark = (remark + "; " + note) if remark else note
        rows.append([x["no"], x["pws"], s["full"], x["ep"], x["kind"], x["when"].strftime("%m/%d/%Y %H:%M"),
                     s["sampler"], text, flow, cl, remark])
    return _w(path, ["state_sample_no", "pws_id", "system_name", "sample_point", "sample_type", "collected",
                     "collector", "sources_in_service", "ep_flow_gpm", "free_chlorine_mg_l", "remarks"], rows)


# ---------------------------------------------------------------------------------------------- the inventory
def _sheet(ws, head, rows, widths):
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE7F0")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append([None if c == "" else c for c in r])
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"


AQUIFERS = ["Sand and gravel", "Cambrian-Ordovician sandstone", "Silurian dolomite", "Glacial outwash"]


def write_inventory(path, S):
    rows = []
    for s in sorted(S, key=lambda x: x["pws"]):
        for w in s["wells"]:
            r = K.rng("facility", w["fid"])
            if w.get("plant"):
                ftype, depth, aq, built = "Surface water treatment plant", "", "", int(r.integers(1958, 1994))
            else:
                ftype = "Well"
                depth = int(10 * round(r.uniform(90, 620) / 10))
                aq = AQUIFERS[int(r.integers(0, 4))]
                built = int(r.integers(1951, 2016)) if not s.get("well_from") else 2025
            sdate = {"Active": "", "Emergency": "2019-03-01", "Inactive": w.get("since", dt.date(2010, 1, 1)).isoformat()
                     if w["status"] == "Inactive" else ""}[w["status"]]
            rows.append([s["pws"], s["full"], w["fid"], w["name"], ftype, w["status"], sdate, w["ep"] or "",
                         w["gpm"], depth, aq, built])
    # the emergency interconnection a listed system opened in 2025
    pax = next(s for s in S if s["name"] == W.PURCHASE[0])
    rows.append([pax["pws"], pax["full"], pax["ic_fid"], pax["ic_name"], "Interconnection",
                 "Emergency", "2011-08-15", "", "", "", "", 2011])
    rows.sort(key=lambda x: (x[0], x[2]))
    wb = Workbook()
    ws = wb.active
    ws.title = "Facilities"
    _sheet(ws, ["pws_id", "system_name", "facility_id", "facility_name", "facility_type", "status", "status_date",
                "entry_point", "rated_capacity_gpm", "well_depth_ft", "aquifer", "year_constructed"], rows,
           [10, 40, 11, 38, 26, 11, 12, 11, 12, 11, 28, 10])
    ns = wb.create_sheet("Notes")
    for line in ["Source facility inventory, community water systems, as of 1 September 2026.",
                 "rated_capacity_gpm: the pump's rated discharge in gallons per minute at the system's design head, "
                 "as the last sanitary survey recorded it. Well pumps in community systems here are constant-speed: a "
                 "well delivers its rated discharge whenever it runs.",
                 "entry_point: the sample point of the entry point to the distribution system the facility feeds.",
                 "status: Active (in routine use), Emergency (kept for emergencies, not in routine use), Inactive "
                 "(disconnected or abandoned); status_date is the date the status took effect."]:
        ns.append([line])
    ns.column_dimensions["A"].width = 120
    wb.save(path)
    return len(rows)


# ---------------------------------------------------------------------------------------------- 2025 operations
def meters(S):
    """The source meters: one on each facility, except where two wells discharge to one header and share it."""
    r = K.rng("meters")
    out, mid = [], 100410
    shared = {(n, a): b for n, (a, b) in W.SHARED}
    shared.update({(n, b): a for n, (a, b) in W.SHARED})
    others = [s for s in S if not s["listed"] and len([w for w in s["wells"] if w["status"] == "Active"]) >= 2
              and not s.get("surface")]
    pick = set(s["name"] for s in [others[i] for i in r.choice(len(others), W.SHARED_OTHER, replace=False)])
    for s in sorted(S, key=lambda x: x["pws"]):
        done = set()
        ws = [w for w in s["wells"] if w["status"] in ("Active", "Emergency")]
        if s["name"] in pick:
            a, b = [w["no"] for w in ws if w["status"] == "Active"][:2]
            shared[(s["name"], a)], shared[(s["name"], b)] = b, a
        for w in ws:
            if w["no"] in done:
                continue
            mid += int(r.integers(1, 30))
            partner = shared.get((s["name"], w["no"]))
            group = [w] + ([x for x in ws if x["no"] == partner] if partner else [])
            for x in group:
                done.add(x["no"])
            m = dict(id="M-%d" % mid, pws=s["pws"], wells=group, size=[2, 3, 4, 6, 8][min(4, int(w["gpm"] // 250))],
                     installed=dt.date(int(r.integers(2004, 2023)), int(r.integers(1, 13)), int(r.integers(1, 28))),
                     removed=None, start=int(r.integers(1_000_000, 900_000_000)))
            for (n, no, when) in W.REPLACED:
                if n == s["name"] and no == w["no"]:
                    m["removed"] = when
                    out.append(m)
                    mid += int(r.integers(1, 30))
                    m = dict(id="M-%d" % mid, pws=s["pws"], wells=group, size=m["size"], installed=when, removed=None,
                             start=0)
            out.append(m)
        if s.get("ic_fid"):
            mid += int(r.integers(1, 30))
            out.append(dict(id="M-%d" % mid, pws=s["pws"], wells=[dict(fid=s["ic_fid"], no=0)], size=6,
                            installed=dt.date(2011, 8, 15), removed=None, start=int(r.integers(10_000, 900_000))))
    return out


def write_meters(path, M):
    rows = []
    for m in M:
        loc = "common discharge header" if len(m["wells"]) > 1 else ("interconnect vault" if m["wells"][0]["no"] == 0
                                                                     else "wellhouse discharge")
        rows.append([m["id"], m["pws"], ";".join(w["fid"] for w in m["wells"]), loc, m["size"], "gallons",
                     m["installed"].isoformat(), m["removed"].isoformat() if m["removed"] else ""])
    return _w(path, ["meter_id", "pws_id", "facility_ids", "location", "size_in", "register_units", "installed",
                     "removed"], rows)


def write_mor(path, S, ops, M):
    """One row a facility a month, as the operator fills the form: the meter's readings at the start and end of
    the month, the gallons through it, and the days the facility ran."""
    by_fid = {}
    for m in M:
        for w in m["wells"]:
            by_fid.setdefault(w["fid"], []).append(m)
    reading = {m["id"]: m["start"] for m in M}
    rows = []
    for s in sorted(S, key=lambda x: x["pws"]):
        ws = [w for w in s["wells"] if w["status"] in ("Active", "Emergency")]
        if not ws:
            continue
        for mth in range(1, 13):
            month_rows = []
            for w in ws:
                ms = by_fid[w["fid"]]
                m = ms[-1]
                days, gal = ops.get((s["name"], w["no"], mth), [0, 0.0])
                meter_gal = sum(ops.get((s["name"], x["no"], mth), [0, 0.0])[1] for x in m["wells"])
                month_rows.append((w, ms, days, int(round(meter_gal))))
            seen = set()
            for w, ms, days, meter_gal in month_rows:
                m = ms[-1]
                first = (m["id"], mth) not in seen
                seen.add((m["id"], mth))
                if len(ms) > 1 and any(x["removed"] and x["removed"].month == mth and x["removed"].year == 2025
                                       for x in ms):
                    old, new = ms[0], ms[-1]
                    frac = (old["removed"].day - 1) / 30.0
                    start = reading[old["id"]]
                    reading[old["id"]] = start + int(round(meter_gal * frac))
                    end = reading.get(new["id"], 0) + int(round(meter_gal * (1 - frac)))
                    reading[new["id"]] = end
                elif len(ms) > 1 and mth < ms[0]["removed"].month:
                    start = reading[ms[0]["id"]]
                    end = start + meter_gal
                    reading[ms[0]["id"]] = end
                else:
                    if first:
                        start = reading[m["id"]]
                        reading[m["id"]] = start + meter_gal
                    else:
                        start = reading[m["id"]] - meter_gal
                    end = start + meter_gal
                rows.append([s["pws"], s["full"], "2025-%02d" % mth, w["fid"], w["name"] if not w.get("plant")
                             else "Treatment plant", start, end, meter_gal, days])
            if s.get("ic_fid"):
                m = by_fid[s["ic_fid"]][-1]
                gal = s["purchase"].get(mth, 0)
                start = reading[m["id"]]
                reading[m["id"]] = start + gal
                rows.append([s["pws"], s["full"], "2025-%02d" % mth, s["ic_fid"], s["ic_name"],
                             start, start + gal, gal, s["purchase_days"].get(mth, 0)])
    return _w(path, ["pws_id", "system_name", "report_month", "facility_id", "facility_name", "meter_start",
                     "meter_end", "volume_gal", "days_in_service"], rows)


# ---------------------------------------------------------------------------------------------- systems and ties
CLOSED = [("BAKKEHAUG MOBILE HOME PARK", "Engvall", 64, "Ground water"),
          ("VILLAGE OF OLD HOVSTAD", "Wrenley", 211, "Ground water"),
          ("MOSSDALE ESTATES", "Larkford", 58, "Ground water"),
          ("STRANDLUND ESTATES", "Hedstrom", 47, "Purchased ground water"),
          ("REDCOMBE SANITARY DISTRICT NO. 3", "Redcombe", 138, "Ground water"),
          ("SANDVIK HEIGHTS WATER ASSOCIATION", "Sandholm", 92, "Ground water")]


def write_register(path, S):
    rows = []
    for s in sorted(S, key=lambda x: x["pws"]):
        r = K.rng("register", s["pws"])
        if s.get("buyer"):
            src = "Purchased ground water" if not s["name"].startswith("Sunrise") else "Purchased surface water"
        elif s.get("surface"):
            src = "Surface water"
        else:
            src = "Ground water"
        conn = int(round(s["pop"] / r.uniform(2.3, 2.7)))
        rows.append([s["pws"], s["full"].upper(), s["county"], s["pop"], src, conn, "Active"])
        if s.get("old_pws"):
            rows.append([s["old_pws"], s["full"].upper(), s["county"], s["old_pop"], src,
                         int(round(s["old_pop"] / r.uniform(2.3, 2.7))), "Inactive"])
    # systems the register still carries after they closed or merged
    r = K.rng("register closed")
    for name, county, pop, src in CLOSED:
        rows.append(["%07d" % int(r.integers(4600000, 4699999)), name, county, pop, src, int(round(pop / 2.5)),
                     "Inactive"])
    rows.sort(key=lambda x: x[0])
    return _w(path, ["pws_id", "system_name", "county", "population_served", "primary_source", "service_connections",
                     "status"], rows)


def write_connections(path, S, conns):
    by = {s["name"]: s for s in S}
    byfull = {s["full"]: s for s in S}
    get = lambda n: by.get(n) or byfull.get(n) or by.get(n.replace("Village of ", "").replace(" Water Utility", ""))
    rows = []
    for c in conns:
        b, s = get(c["buyer"]), get(c["seller"])
        rows.append([b.get("old_pws", b["pws"]), b["full"], s.get("old_pws", s["pws"]), s["full"], c["kind"],
                     c["start"].isoformat(), c["end"].isoformat() if c["end"] else ""])
    pax = by[W.PURCHASE[0]]
    hr = next(s for s in S if s.get("is_regional"))
    rows.append([pax["pws"], pax["full"], hr["pws"], hr["full"], "Emergency", "2011-08-15", ""])
    rows.sort(key=lambda x: x[0])
    return _w(path, ["buyer_pws_id", "buyer_name", "seller_pws_id", "seller_name", "connection_type",
                     "in_service_from", "in_service_to"], rows)


def write_labs(path):
    rows = [[c, n, city, m, cert, "PFAS (%s); %s" % (m, "inorganics, VOCs" if c == "SEHL" else "SOCs")]
            for c, n, city, m, cert in LABS]
    return _w(path, ["lab_code", "laboratory", "city", "pfas_method", "state_certification_no", "certified_for"], rows)


# ---------------------------------------------------------------------------------------------- kept beside them
UCMR_MRL = {"PFOA": 0.004, "PFOS": 0.004, "PFHxS": 0.003, "PFNA": 0.004, "PFBS": 0.003, "HFPO-DA": 0.005,
            "PFBA": 0.005, "PFPeA": 0.003, "PFHxA": 0.001, "PFHpA": 0.003, "PFDA": 0.003, "PFUnA": 0.002,
            "PFDoA": 0.003, "PFPeS": 0.004, "PFHpS": 0.003, "4:2 FTS": 0.003, "6:2 FTS": 0.005, "8:2 FTS": 0.005,
            "ADONA": 0.003, "PFMPA": 0.004, "PFMBA": 0.003, "NFDHA": 0.02, "9Cl-PF3ONS": 0.002,
            "11Cl-PF3OUdS": 0.005, "PFEESA": 0.003, "PFTrDA": 0.007, "PFTA": 0.008, "NMeFOSAA": 0.006,
            "NEtFOSAA": 0.005}


def write_survey(path, S):
    """The federal survey's two 2023 sampling events at the entry points of the state's larger systems and a sample of
    the smaller ones, as EPA published them: micrograms per litre, the survey's own reporting levels, its own facility
    and sample point numbers."""
    rows = []
    r0 = K.rng("survey")
    for s in sorted(S, key=lambda x: x["pws"]):
        if not s["eps"] or s.get("well_from"):
            continue
        if s["pop"] < 3300 and r0.random() > 0.35:
            continue
        size = "L" if s["pop"] > 10000 else "S"
        for e, ep in enumerate(s["eps"]):
            fac = "%05d" % int(r0.integers(10000, 99999))
            r = K.rng("survey", s["pws"], ep["code"])
            first = dt.date(2023, int(r.integers(2, 5)), int(r.integers(1, 28)))
            for ev, day in (("SE1", first), ("SE2", first + dt.timedelta(days=int(r.integers(150, 200))))):
                t = dt.datetime.combine(day, dt.time(10))
                on = W.in_service(r, ep, t)
                rate = {w["no"]: w["gpm"] for w in ep["wells"]}
                true = W.blend(ep, on, {n: rate[n] for n in on})
                sid = "%d" % int(r.integers(100000000, 999999999))
                for a in sorted(UCMR_MRL):
                    v = true.get(a, 0.0) * math.exp(r.normal(0, 0.08)) / 1000.0
                    mrl = UCMR_MRL[a]
                    sign = "=" if v >= mrl else "<"
                    rows.append([s["pws"], s["full"].upper(), size, fac, "EP %s" % ep["code"][2:],
                                 "GW" if not s.get("surface") else "SW", "EPTDS%02d" % (e + 1),
                                 "ENTRY POINT %d" % (e + 1), "EP", day.strftime("%m/%d/%Y"), sid, a, "%.3f" % mrl,
                                 "533" if a not in ("PFTrDA", "PFTA", "NMeFOSAA", "NEtFOSAA") else "537.1", sign,
                                 "%.4f" % v if sign == "=" else "", ev])
    return _w(path, ["PWSID", "PWSName", "Size", "FacilityID", "FacilityName", "FacilityWaterType", "SamplePointID",
                     "SamplePointName", "SamplePointType", "CollectionDate", "SampleID", "Contaminant", "MRL",
                     "MethodID", "AnalyticalResultsSign", "AnalyticalResultValue", "SampleEventCode"], rows)


TOWNS = {"Granlund": ["Town of Nyquist", "Town of Tolvik", "Town of Bakkehaug", "Town of Hagestad", "Town of Ingvoldstad"],
         "Engvall": ["Town of Kvilund", "Town of Urdahl", "Town of Torvang", "Town of Lokberg", "Town of Hovstad"],
         "Hedstrom": ["Town of Svalstad", "Town of Strandlund", "Town of Fjellby", "Town of Ulstad", "Town of Granvold"],
         "Sandholm": ["Town of Sletstad", "Town of Fosvik", "Town of Nesdahl", "Town of Vesstad", "Town of Jorrud"]}


def write_private(path):
    """The department's private well program: homeowners' samples sent to the state laboratory, 2024 to mid-2026."""
    r = K.rng("private")
    rows = []
    n = 0
    for county, towns in TOWNS.items():
        for town in towns:
            hot = r.random() < 0.3
            for i in range(int(r.integers(14, 32))):
                n += 1
                day = dt.date(2024, 1, 2) + dt.timedelta(days=int(r.integers(0, 900)))
                level = math.exp(r.normal(1.2 if hot else -0.6, 1.1))
                vals = {"PFOA": level * r.uniform(0.3, 0.9), "PFOS": level * r.uniform(0.4, 1.3)}
                vals["PFHxS"] = vals["PFOS"] * r.uniform(0.2, 0.7)
                vals["PFNA"] = vals["PFOA"] * r.uniform(0.0, 0.1)
                vals["HFPO-DA"] = 0.0
                vals["PFBS"] = vals["PFOS"] * r.uniform(0.1, 0.4)
                out = [fmt(v) if v >= 1.0 else "<1.0" for v in vals.values()]
                tot = sum(v for k, v in vals.items() if k in ("PFOA", "PFOS") and v >= 1.0)
                rows.append(["PW-%05d" % (40000 + n * 7 + int(r.integers(0, 6))), county, town, day.isoformat(),
                             "SEHL"] + out + ["%.1f" % tot, "Y" if tot > 20 else "N"])
    rows.sort(key=lambda x: x[3])
    return _w(path, ["sample_id", "county", "municipality", "collected", "lab_code", "PFOA_ng_l", "PFOS_ng_l",
                     "PFHxS_ng_l", "PFNA_ng_l", "HFPO-DA_ng_l", "PFBS_ng_l", "PFOA_plus_PFOS_ng_l",
                     "above_state_advisory_20"], rows)
