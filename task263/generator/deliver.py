"""The goldens: the 2027 grant workbook (the list with each system's grant and the walk, the wells' own water, the
county figures) and the chart of each active well's PFOA and PFOS against the line, both from analysis.py on the
shipped pack; and generator/figures.json, the figures the write-up states."""
from __future__ import annotations

import json
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt                       # noqa: E402
from openpyxl import Workbook                          # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

import analysis as A                                   # noqa: E402
import common as K                                     # noqa: E402
import postprocess as PP                               # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
COUNTIES = ("Granlund", "Engvall", "Hedstrom", "Sandholm")


def results(target):
    D = A.load(target)
    G = A.golden(D)
    rows = A.window(D)
    n_ep = {}
    for x in rows:
        n_ep[(x["pws"], x["ep"])] = n_ep.get((x["pws"], x["ep"]), 0) + 1
    return D, G, A.ask_a(D, target), A.ask_b(D, target), n_ep


def _head(ws, head, widths):
    ws.append(head)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E3EAF2")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w


def workbook(path, D, G, ask_a, ask_b, n_ep):
    wb = Workbook()
    ws = wb.active
    ws.title = "2027 grants"
    _head(ws, ["Rank", "PWS ID", "System", "County", "Wells over 4.0 ng/L", "Rated capacity of those wells (gpm)",
               "MGD", "Grant ($)", "Funded in 2027", "Set-aside left after ($)"], [6, 10, 38, 9, 22, 14, 8, 12, 14, 14])
    left = K.FUND
    funded = set(G["funded"])
    for p in D["plan"]:
        over = sorted(G["over"].get(p["pws"], set()))
        gpm = sum(D["wells"][(p["pws"], ep)][n] for ep, n in over)
        wells = ", ".join("Well %d" % n for ep, n in over)
        g = G["grants"][p["pws"]]
        if p["pws"] in funded:
            left -= g
            ws.append([p["rank"], p["pws"], p["name"], p["county"], wells, gpm, round(gpm * 1440 / 1e6, 3), g, "Yes", left])
        else:
            ws.append([p["rank"], p["pws"], p["name"], p["county"], wells, gpm, round(gpm * 1440 / 1e6, 3), g,
                       "No (2028 list)", None])
        ws.cell(ws.max_row, 8).number_format = "#,##0"
        ws.cell(ws.max_row, 10).number_format = "#,##0"
    stop = next(p for p in D["plan"] if p["pws"] == G["stop"])
    ws.append([])
    for label, v in (("Systems the $36,000,000 reaches", len(G["funded"])),
                     ("First system it cannot reach", "%s (grant $%s)" % (stop["name"], "{:,}".format(G["grants"][stop["pws"]]))),
                     ("Set-aside left ($)", G["left"]),
                     ("Grants made ($)", K.FUND - G["left"])):
        ws.append([None, None, label, None, None, None, None, v])
        ws.cell(ws.max_row, 3).font = Font(bold=True)
        if isinstance(v, int):
            ws.cell(ws.max_row, 8).number_format = "#,##0"
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Wells")
    _head(ws, ["System", "Entry point", "Well", "Rated capacity (gpm)", "Own PFOA (ng/L)", "Own PFOS (ng/L)",
               "Over 4.0 ng/L", "Samples at the entry point"], [38, 10, 8, 12, 12, 12, 10, 12])
    for p in D["plan"]:
        for (pws, ep), lv in sorted(G["levels"].items()):
            if pws != p["pws"]:
                continue
            for n in sorted(lv):
                pfoa, pfos = lv[n]
                ws.append([p["name"], ep, "Well %d" % n, D["wells"][(pws, ep)][n], round(pfoa, 1), round(pfos, 1),
                           "Yes" if max(pfoa, pfos) > K.LINE else "No", n_ep[(pws, ep)]])
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Counties")
    _head(ws, ["County", "Water pumped by the listed systems' wells, 2025 (million gallons)",
               "People who get their water from those wells"], [12, 26, 24])
    for c in COUNTIES:
        ws.append([c, ask_a[c], ask_b[c]])
        ws.cell(ws.max_row, 2).number_format = "#,##0.0"
        ws.cell(ws.max_row, 3).number_format = "#,##0"
    ws.append(["Total", round(sum(ask_a[c] for c in COUNTIES), 1), sum(ask_b[c] for c in COUNTIES)])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.cell(ws.max_row, 2).number_format = "#,##0.0"
    ws.cell(ws.max_row, 3).number_format = "#,##0"
    ws.append([])
    ws.append(["Water: 2025 monthly operating reports, each meter once a month. People: the system register's "
               "populations, 1 September 2026."])

    ws = wb.create_sheet("Notes")
    ws.column_dimensions["A"].width = 120
    blank = sum(1 for x in D["samples"] if x["listed"] and x["on"] is None and x["kind"] in ("RT", "CO")
                and x["date"].year in K.WINDOW)
    words = {1: "one", 2: "two", 3: "three", 4: "four", 5: "five", 6: "six", 7: "seven", 8: "eight", 9: "nine"}
    for line in NOTES:
        ws.append([line.replace("{blank}", words.get(blank, str(blank)))])
        ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


NOTES = [
    "Prepared for the board papers, 18 November 2026. Sources: the department's 2 September 2026 data, the draft plan's section 6.",
    "Wells. Every entry-point sample is the water of the wells pumping when it was drawn, mixed in proportion to their "
    "rated capacities (the entry-point flow on each record matches the sum to within a few per cent). For each entry "
    "point the 2024 and 2025 routine and confirmation samples (latest report revision; blanks and duplicates left out; "
    "non-detects at half the reporting limit; the {blank} samples with no sources recorded left out) were solved together "
    "by least squares for each well's own PFOA and PFOS.",
    "The same wells come out over 4.0 ng/L with non-detects at zero or at the reporting limit, by non-negative least "
    "squares, with the entry-point flow or equal shares as the mix, on routine samples alone, and with the 2026 samples "
    "added; no well's estimate sits close enough to the line for the choice to matter.",
    "Grants: $2.4 million per MGD of the rated capacity of the active wells over 4.0 ng/L, capped at $6 million, in the "
    "order of Table 6-1 until the remaining set-aside does not cover the next grant.",
    "Counties. Water: the 2025 monthly operating reports for the listed systems' wells, each meter counted once a month "
    "(Lindvik Lake, Halvorsen, Skogen and Elvholm Falls each have two wells on one header meter, written on both "
    "wells' lines); Ackerholt Creek's September purchase through its emergency interconnect is not well water and is "
    "left out; Kesselring's emergency well (three days in July) is in.",
    "People: each listed system's own customers plus the customers of every system that buys all its water from it, "
    "directly or through another buyer, on the connections in service (Solberg Lake's emergency tie and Hagestad "
    "Corners' purchase, closed 30 June 2025, are left out), counted in the county where the buyer is.",
]


def chart(path, D, G):
    labels, pfoa, pfos, over = [], [], [], []
    for p in D["plan"]:
        for (pws, ep), lv in sorted(G["levels"].items()):
            if pws != p["pws"]:
                continue
            for n in sorted(lv):
                labels.append("%s, Well %d" % (p["name"].replace(" Water Utility", "").replace("Village of ", "")
                                              .replace(" Municipal Water", "").replace(" Water and Light", "")
                                              .replace(" Sanitary District No. 1", "").replace(" Sanitary District No. 2", "")
                                              .replace(" Sanitary District", "").replace("City of ", "")
                                              .replace(" Water Department", "").replace(" Waterworks", "")
                                              .replace("Town of ", ""), n))
                pfoa.append(max(lv[n][0], 0.2))
                pfos.append(max(lv[n][1], 0.2))
                over.append(max(lv[n]) > K.LINE)
    n = len(labels)
    fig, ax = plt.subplots(figsize=(8.6, 11.6), dpi=110)
    y = list(range(n))[::-1]
    for yi, a, b, o in zip(y, pfoa, pfos, over):
        ax.plot([min(a, b), max(a, b)], [yi, yi], color="#c9d2da", lw=1, zorder=1)
    ax.scatter(pfoa, y, s=26, color="#1f5f9e", label="PFOA", zorder=3)
    ax.scatter(pfos, y, s=26, marker="D", color="#d0782a", label="PFOS", zorder=3)
    ax.axvline(K.LINE, color="#b3261e", lw=1.3, ls="--", zorder=2)
    ax.text(K.LINE * 1.04, n - 0.3, "4.0 ng/L", color="#b3261e", fontsize=8, va="bottom")
    ax.set_xscale("log")
    ax.set_xlim(0.2, 50)
    ax.set_xticks([0.3, 1, 2, 4, 10, 20, 40])
    ax.set_xticklabels(["0.3", "1", "2", "4", "10", "20", "40"])
    ax.set_yticks(y)
    ax.set_yticklabels(labels, fontsize=6.6)
    for lab, o in zip(ax.get_yticklabels(), over):
        if o:
            lab.set_fontweight("bold")
            lab.set_color("#7a1712")
    ax.set_ylim(-0.8, n - 0.2)
    ax.set_xlabel("Own water, 2024 to 2025 average (ng/L, log scale)", fontsize=8.5)
    ax.set_title("%d of the listed systems' %d active wells carry more than 4.0 ng/L of PFOA or PFOS\n"
                 "own water, 2024 to 2025, solved from the entry-point samples; wells over the line in bold red; "
                 "systems in list order" % (sum(over), n), fontsize=9.2, loc="left")
    ax.grid(axis="x", color="#eceff2", lw=0.6)
    ax.legend(loc="lower right", fontsize=8, frameon=False)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    fig.tight_layout()
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)


def figures(D, G, ask_a, ask_b):
    name = {p["pws"]: p["name"] for p in D["plan"]}
    rows = A.window(D)
    blank = sum(1 for x in D["samples"] if x["listed"] and x["on"] is None and x["kind"] in ("RT", "CO")
                and x["date"].year in K.WINDOW)
    F = dict(call=dict(systems_reached=len(G["funded"]), first_not_reached=name[G["stop"]], left=G["left"],
                       grants_made=K.FUND - G["left"], stop_grant=G["grants"][G["stop"]]),
             counts=dict(samples=len(rows), entry_points=len({(x["pws"], x["ep"]) for x in rows}), blank=blank,
                         wells_active=sum(len(lv) for lv in G["levels"].values()),
                         wells_over=sum(len(v) for v in G["over"].values())),
             mgd_rule=[1440, 1000000], plan=dict(fund_millions=36, rate_millions=2.4, cap_millions=6, line=4.0,
                                                    sections=["6.3", "6.4"]),
             grants={name[p["pws"]]: G["grants"][p["pws"]] for p in D["plan"]},
             funded=[name[x] for x in G["funded"]],
             wells_over={name[pws]: sorted("Well %d" % n for ep, n in v) for pws, v in G["over"].items()},
             levels={"%s %s Well %d" % (name[pws], ep, n): [round(a, 1), round(b, 1)]
                     for (pws, ep), lv in G["levels"].items() if pws in name for n, (a, b) in lv.items()},
             ask_a=ask_a, ask_b=ask_b)
    return F


def write(target, golden, T=None):
    os.makedirs(golden, exist_ok=True)
    D, G, aa, ab, n_ep = results(target)
    xp = os.path.join(golden, K.DELIVERABLES[0])
    workbook(xp, D, G, aa, ab, n_ep)
    PP.COMPANY = "Water Infrastructure Finance Authority"
    PP.normalize_ooxml(xp, K.PEOPLE["manager"], K.PEOPLE["manager"], "2026-10-07T15:12:40Z", "2026-10-09T19:47:03Z",
                       excel=True)
    chart(os.path.join(golden, K.DELIVERABLES[1]), D, G)
    F = figures(D, G, aa, ab)
    json.dump(F, open(os.path.join(HERE, "figures.json"), "w"), indent=1, sort_keys=True)
    meta = {"distractor_files": ["target/" + K.FILES[k] for k in K.DISTRACTORS]}
    json.dump(meta, open(os.path.join(os.path.dirname(target), "metadata.json"), "w"), indent=2)
    return F
