"""Reads the cut pack back: each group's paid medical cost and member-months under each route, and the pharmacy
summary by drug class."""
import csv
import datetime as dt
import os

from openpyxl import load_workbook

import world as W


def find(folder, stem):
    for fn in sorted(os.listdir(folder)):
        if fn.startswith(stem):
            return os.path.join(folder, fn)
    raise FileNotFoundError(stem)


def _rows(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def groups(target):
    ws = load_workbook(find(target, "provider_group_register"), read_only=True).active
    return sorted(r[0] for r in ws.iter_rows(min_row=5, values_only=True) if r and r[0])


MONTHS = {m.strftime("%Y-%m") for m in W.MONTHS}


def roster(target):
    """{(member, month)}: the roster's member-months, and {member: group}."""
    on, grp = set(), {}
    for r in _rows(find(target, "enrollment_roster")):
        if r["roster_month"] in MONTHS:
            on.add((r["member_id"], r["roster_month"]))
            grp[r["member_id"]] = r["pcp_group_id"]
    return on, grp


def crosswalk(target):
    """{state Medicaid id: plan member id}, from the roster."""
    return {r["medicaid_id"]: r["member_id"] for r in _rows(find(target, "enrollment_roster"))}


def member_lines(target):
    """{(member, coverage month): signed member count} over the remittance's member-level lines."""
    net, cw = {}, crosswalk(target)
    for r in _rows(find(target, "capitation_remittance")):
        if r["medicaid_id"] and r["coverage_month"] in MONTHS:
            k = (cw[r["medicaid_id"]], r["coverage_month"])
            net[k] = net.get(k, 0) + int(r["member_count"])
    return net


def claims(target):
    """[(member, service month, paid)] over the paid medical claims in the period, paid by the cut-off."""
    out = []
    for r in _rows(find(target, "medical_claims_paid")):
        if r["service_from"][:7] in MONTHS and r["paid_date"] <= W.PAID_BY.isoformat():
            out.append((r["member_id"], r["service_from"][:7], float(r["paid_amount"])))
    return out


def exposure(target, route="golden"):
    """{group: member-months}. golden: the months each member was covered, the roster's months less those the state
    recouped and plus those it paid capitation for off the roster; roster: the roster's months; retro: plus the
    months paid off the roster only; recoup: less the recouped months only; orphan: the roster's months plus
    off-roster months with a paid claim."""
    on, grp = roster(target)
    net = member_lines(target)
    months = set(on)
    if route in ("golden", "retro"):
        months |= {k for k, v in net.items() if v > 0 and k not in on}
    if route in ("golden", "recoup"):
        months -= {k for k, v in net.items() if v < 0 and k in on}
    if route == "orphan":
        months |= {(m, s) for m, s, _ in claims(target)}
    out = {}
    for m, s in months:
        if m in grp:
            out[grp[m]] = out.get(grp[m], 0) + 1
    return out


def cost(target):
    _, grp = roster(target)
    out = {}
    for m, s, p in claims(target):
        if m in grp:
            out[grp[m]] = out.get(grp[m], 0.0) + p
    return out


def pmpm(target, route="golden"):
    c, e = cost(target), exposure(target, route)
    return {g: c.get(g, 0.0) / e[g] for g in groups(target)}


def classes(target):
    ws = load_workbook(find(target, "drug_class_crosswalk"), read_only=True).active
    return {r[0]: r[1] for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]}


def pharmacy(target, route="golden"):
    """{class: (paid dollars, fills)}. golden: each claim number at its latest line, fills dated in the period, a
    fill being a claim whose latest line paid something; all: every line counted; filterfirst: lines filtered to the
    period and class first, then each claim number at its latest line among them; first: each claim number at its
    first line."""
    ndc = classes(target)
    lines = _rows(find(target, "pharmacy_claims"))
    a, b = W.MONTHS[0].isoformat(), "2026-06-30"
    out = {c: [0.0, 0] for c in sorted(set(ndc.values()))}
    if route == "all":
        for r in lines:
            if a <= r["fill_date"] <= b:
                out[ndc[r["ndc"]]][0] += float(r["paid_amount"])
                out[ndc[r["ndc"]]][1] += 1
        return {c: (round(v[0], 2), v[1]) for c, v in out.items()}
    if route == "filterfirst":
        lines = [r for r in lines if a <= r["fill_date"] <= b]
        key = lambda r: (r["claim_number"], ndc[r["ndc"]])
    else:
        key = lambda r: r["claim_number"]
    pick = {}
    for r in lines:
        k = key(r)
        if k not in pick or (r["processed_at"] > pick[k]["processed_at"]) == (route != "first"):
            pick[k] = r
    for r in pick.values():
        if a <= r["fill_date"] <= b:
            p = float(r["paid_amount"])
            out[ndc[r["ndc"]]][0] += p
            out[ndc[r["ndc"]]][1] += 1 if p > 0 else 0
    return {c: (round(v[0], 2), v[1]) for c, v in out.items()}
