"""The golden deliverables, computed from the shipped files only: rof_fy2027_region_poverty.csv (each planning region's
residents in poverty, 2020 to 2024, and its 2020-2024 average) and rof_fy2027_allocation_memo.docx (the State Data
Center's memo to the Office of Budget and Policy: the nine FY2027 allocations with a bar chart, and the transition
figures); and figures.json, every figure the write-up states."""
import csv
import io
import json
import os
import sys

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import askpaths as AP   # noqa: E402
import bundle as B      # noqa: E402
import world as W       # noqa: E402

LEGACY = (2020, 2021, 2022)
YEARS = (2020, 2021, 2022, 2023, 2024)
FUND = 12_000_000
BLUE = "#2a78d6"


def r100(v):
    return int(round(v / 100.0)) * 100


def r10k(v):
    return int(round(v / 10000.0)) * 10000


def fmt(n):
    return "{:,}".format(int(n))


def series(P, est):
    """Each region's residents in poverty, 2020 to 2024: the former counties' estimates divided among their towns by
    `est(town, year)` for 2020 to 2022 (rule 2.2), the regions' own estimates for 2023 and 2024 (rule 2.1)."""
    out = {g: {} for g in P.regions}
    for y in LEGACY:
        for g in P.regions:
            out[g][y] = 0.0
        for c in P.counties:
            ts = [t for t in P.towns if P.county_of[t] == c]
            v = {t: est(t, y) for t in ts}
            den = sum(v.values())
            for t in ts:
                out[P.region_of[t]][y] += P.sdc[(c, y)] * v[t] / den
    for y in (2023, 2024):
        for g in P.regions:
            out[g][y] = float(P.sdc[(g, y)])
    return out


def allocate(ser, regions):
    avg = {g: float(np.mean([ser[g][y] for y in YEARS])) for g in regions}
    tot = sum(avg.values())
    return avg, {g: FUND * avg[g] / tot for g in regions}, tot


def compute(target):
    P = B.Pack(target)
    sy = {(t, y): P.snap_year(t, y) for t in P.towns for y in YEARS}
    mean5 = {t: float(np.mean([sy[(t, y)] for y in YEARS])) for t in P.towns}
    ratio = {t: P.below[t] / mean5[t] for t in P.towns}
    routes = {
        "golden": lambda t, y: sy[(t, y)] * ratio[t],
        "release_shares": lambda t, y: float(P.below[t]),
        "snap_shares": lambda t, y: sy[(t, y)],
        "middle_year": lambda t, y: P.below[t] * sy[(t, y)] / sy[(t, 2022)],
    }
    out = {}
    for k, est in routes.items():
        ser = series(P, est)
        avg, amt, tot = allocate(ser, P.regions)
        out[k] = dict(series=ser, avg=avg, amount=amt, total=tot)
    ext, aud, conv, bal = P.asks_inputs()
    asks = AP.figures(ext, aud, conv, bal)
    return dict(P=P, routes=out, ratio=ratio, asks=asks, balances=bal, ext_rows=len(ext))


def figures(X):
    P, G = X["P"], X["routes"]["golden"]
    regs = P.regions
    a = X["asks"]
    f = {
        # the committed call for score_reader.py: the nine allocations as one list in region-code order; the
        # per-region amounts and the conformed series are the call's inputs (keys carrying "rows" and "population")
        "call": {
            "allocation": [r10k(G["amount"][g]) for g in regs],
            "allocation_rows": {g: r10k(G["amount"][g]) for g in regs},
            "population_by_year": {g: {str(y): r100(G["series"][g][y]) for y in LEGACY} for g in regs},
            "population_average": {g: r100(G["avg"][g]) for g in regs},
        },
        "regions": [{"region": g, "code": [r["region_fips"] for r in P.codes if r["planning_region"] == g][0],
                     "council": P.council_of[g]} for g in regs],
        "amounts": {g: r10k(G["amount"][g]) for g in regs},
        "amounts_total_rounded": sum(r10k(G["amount"][g]) for g in regs),
        "fund": FUND,
        "average": {g: r100(G["avg"][g]) for g in regs},
        "average_sum": r100(G["total"]),
        "cells": {g: {str(y): r100(G["series"][g][y]) for y in LEGACY} for g in regs},
        "published": {g: {str(y): r100(G["series"][g][y]) for y in (2023, 2024)} for g in regs},
        "share_pct": {g: round(100.0 * G["avg"][g] / G["total"], 1) for g in regs},
        "rivals": {k: {g: r10k(X["routes"][k]["amount"][g]) for g in regs}
                   for k in ("release_shares", "snap_shares", "middle_year")},
        "rivals_off": {k: sum(r10k(X["routes"][k]["amount"][g]) != r10k(G["amount"][g]) for g in regs)
                       for k in ("release_shares", "snap_shares", "middle_year")},
        "ratio_range": [round(min(X["ratio"].values()), 2), round(max(X["ratio"].values()), 2)],
        "county_totals": {str(y): sum(P.sdc[(c, y)] for c in P.counties) for y in LEGACY},
        "release_months": 60,
        "rounding": {"amounts": 10000, "residents": 100, "dollars": 100},
        "rule_sections": {"fund": "1.1", "latest": "2.1", "division": "2.2", "average": "2.3", "allocation": "2.4",
                          "balances": "3.1", "energy": "3.2"},
        "asks": {
            "carry": {g: int(round(a["carry"][g] / 100.0)) * 100 for g in regs},
            "people": {g: int(a["people"][g]) for g in regs},
            "energy": {g: int(round(a["energy"][g] / 100.0)) * 100 for g in regs},
            "unexpended": {c: X["balances"][c]["unspent"] for c in P.counties},
            "unexpended_total": sum(X["balances"][c]["unspent"] for c in P.counties),
        },
    }
    return f


CSV_COLS = ["region_code", "planning_region", "council_of_governments", "residents_in_poverty_2020",
            "residents_in_poverty_2021", "residents_in_poverty_2022", "residents_in_poverty_2023",
            "residents_in_poverty_2024", "average_2020_2024"]


def write_csv(F, path):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(CSV_COLS)
        for r in F["regions"]:
            g = r["region"]
            c, p = F["cells"][g], F["published"][g]
            w.writerow([r["code"], g, r["council"], c["2020"], c["2021"], c["2022"], p["2023"], p["2024"],
                        F["average"][g]])


def chart_png(F):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.family": "DejaVu Sans", "svg.hashsalt": "task474", "font.size": 9})
    names = [r["council"].replace(" Council of Governments", "") for r in F["regions"]]
    vals = [F["amounts"][r["region"]] / 1e6 for r in F["regions"]]
    fig, ax = plt.subplots(figsize=(6.4, 3.6), dpi=200)
    y = np.arange(len(names))[::-1]
    ax.barh(y, vals, height=0.56, color=BLUE, zorder=2)
    for yi, v in zip(y, vals):
        ax.text(v + 0.03, yi, "$%s" % fmt(v * 1e6), va="center", ha="left", fontsize=8, color="#1b1d20")
    ax.set_yticks(y)
    ax.set_yticklabels(names, fontsize=8.5, color="#1b1d20")
    ax.set_xlim(0, max(vals) * 1.22)
    ax.set_xlabel("FY2027 allocation ($ millions)", fontsize=8.5, color="#52514e")
    ax.xaxis.grid(True, color="#e3e2de", linewidth=0.8, zorder=0)
    ax.tick_params(axis="x", colors="#52514e", labelsize=8)
    ax.tick_params(axis="y", length=0)
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.spines["bottom"].set_color("#b9b8b2")
    top = sorted(F["regions"], key=lambda r: -F["amounts"][r["region"]])[:2]
    ax.set_title("%s and %s receive the largest FY2027 allocations" % (top[0]["region"], top[1]["region"]),
                 fontsize=10, loc="left", color="#1b1d20")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def write_docx(F, path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1.0)
        s.top_margin = s.bottom_margin = Inches(0.9)

    def para(text, bold=False, size=None, space=6, italic=False, color=None):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        p.paragraph_format.space_after = Pt(space)
        return p

    def table(rows, widths, right_from=1):
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        t.autofit = False
        for j, wd in enumerate(widths):
            t.columns[j].width = Inches(wd)
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                c = t.cell(i, j)
                c.width = Inches(widths[j])
                c.text = ""
                run = c.paragraphs[0].add_run(str(v))
                run.font.size = Pt(9)
                run.bold = (i == 0) or (i == len(rows) - 1 and str(row[0]).startswith("Total"))
                if j >= right_from and i > 0:
                    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    foot = doc.sections[0].footer.paragraphs[0]
    fr = foot.add_run("%s  |  Regional Opportunity Fund, FY2027 allocation  |  Page " % W.SDC)
    fr.font.size = Pt(8)
    run = foot.add_run()
    run.font.size = Pt(8)
    for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "end")):
        el = OxmlElement(tag)
        if attr:
            el.set(qn("w:fldCharType"), attr)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = "PAGE"
        run._r.append(el)
    para("STATE DATA CENTER", bold=True, size=9, space=0, color="1F4E79")
    para("Memorandum", bold=True, size=15, space=8)
    for k, v in (("To", "%s, Grants Manager, %s" % (W.PEOPLE["grants"], W.OBP)),
                 ("From", "%s, Director, and %s, Demographer, %s" % (W.PEOPLE["director"], W.PEOPLE["demographer"],
                                                                     W.SDC)),
                 ("Date", "October 16, 2026"),
                 ("Ref", "SDC 26-114"),
                 ("Subject", "Regional Opportunity Fund: FY2027 allocations to the nine councils and the FY2026 "
                             "transition figures")):
        p = doc.add_paragraph()
        r1 = p.add_run("%s: " % k)
        r1.bold = True
        p.add_run(v)
        p.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    regs = [r["region"] for r in F["regions"]]
    para("FY2027 allocation", bold=True, size=11.5, space=4)
    para("Under the allocation rule (sections 1 and 2), the $12,000,000 is split in proportion to each planning "
         "region's average residents in poverty over 2020 to 2024. The allocations are:")
    rows = [["Council of governments", "Average residents in poverty, 2020-2024", "Share", "FY2027 allocation"]]
    for r in F["regions"]:
        g = r["region"]
        rows.append([r["council"], fmt(F["average"][g]), "%.1f%%" % F["share_pct"][g], "$" + fmt(F["amounts"][g])])
    rows.append(["Total", fmt(F["average_sum"]), "100.0%", "$" + fmt(F["amounts_total_rounded"])])
    table(rows, [2.9, 1.5, 0.7, 1.3])
    para("Each allocation is computed from the unrounded averages and rounded to the nearest $10,000 (rule 2.4), so "
         "the nine rounded amounts total $%s; the unrounded allocations total $%s. Averages are rounded to the "
         "nearest hundred. Sources: %s estimates as last published (December 2022 to December 2025 releases); U.S. "
         "Census Bureau, 2020-2024 American Community Survey, table B17001; Department of Social Services, SNAP "
         "participation by town." % (fmt(F["amounts_total_rounded"]), fmt(F["fund"]), W.SDC), size=9, italic=True)
    doc.add_picture(io.BytesIO(chart_png(F)), width=Inches(6.2))
    para("How the 2020 to 2022 figures were built", bold=True, size=11.5, space=4)
    para("Our estimates for 2020, 2021 and 2022 exist only for the eight former counties, so each former county's "
         "estimate (as last published) was divided among its towns by each town's residents in poverty in that "
         "year, as rule 2.2 requires. The Bureau's 2020-2024 five-year table is the only count of the poor by town, "
         "but it is an average over all five years and already reflects the 2021 moves out of Ormsby and Wexport "
         "for most of them, so its shares are not where the poor lived in 2020 and 2021. Each town's figure for a year is "
         "therefore its average monthly SNAP recipients that year, scaled by the town's own ratio of the five-year "
         "survey count to its average recipients over 2020 to 2024. The ratio runs from %.2f to %.2f across the "
         "towns, because SNAP participation among the poor differs from town to town. The 2023 and 2024 figures are "
         "our planning region estimates (2023 revised, 2024 first)." % tuple(F["ratio_range"]))
    rows = [["Planning region", "2020", "2021", "2022", "2023", "2024"]]
    for g in regs:
        c, p = F["cells"][g], F["published"][g]
        rows.append([g, fmt(c["2020"]), fmt(c["2021"]), fmt(c["2022"]), fmt(p["2023"]), fmt(p["2024"])])
    table(rows, [1.9, 0.9, 0.9, 0.9, 0.9, 0.9])
    para("Residents in poverty, nearest hundred. The full series is in rof_fy2027_region_poverty.csv. The 2024 "
         "figures are first estimates; under our methods note they are revised once, in December 2026, after this "
         "allocation, and the rule takes them as last published.", size=9, italic=True)
    para("Transition from the former county agencies", bold=True, size=11.5, space=4)
    para("Under rule 3.1 the agencies' unexpended FY2026 balances ($%s in all) pass to the councils in proportion to "
         "the people each agency served in FY2026 while they lived in each council's region; under rule 3.2 each "
         "council's FY2027 energy assistance plan starts from the energy assistance paid in FY2026 to households in "
         "its towns." % fmt(F["asks"]["unexpended_total"]))
    rows = [["Council of governments", "Share of unexpended FY2026 balances", "People counted",
             "Energy assistance paid, FY2026"]]
    for r in F["regions"]:
        g = r["region"]
        rows.append([r["council"], "$" + fmt(F["asks"]["carry"][g]), fmt(F["asks"]["people"][g]),
                     "$" + fmt(F["asks"]["energy"][g])])
    table(rows, [2.9, 1.3, 1.0, 1.2])
    para("Balances and energy assistance to the nearest $100. The people counts come from the grants data system's "
         "extract of September 30, 2026: Wexmoor's April 2026 upload, received twice, is counted once; services are "
         "kept by service date (July 1, 2025 to June 30, 2026); Corbin's earlier client numbers are matched to the "
         "numbers it issued on January 1, 2026; and each client is placed in the town the profile audit trail gives "
         "for the service date, since the extract shows each record's current town. A person served in two regions "
         "counts in each. An energy assistance payment appears on the row of every member of the household and is "
         "counted once. Sources: grants data system extract and profile audit trail, September 30, 2026; Corbin "
         "Community Action client number conversion, January 1, 2026; FY2026 closeout, October 2, 2026.", size=9,
         italic=True)
    doc.save(path)


def main(target, golden, figures_path):
    import postprocess as PP
    os.makedirs(golden, exist_ok=True)
    X = compute(target)
    F = figures(X)
    write_csv(F, os.path.join(golden, W.DELIVERABLES[0]))
    write_docx(F, os.path.join(golden, W.DELIVERABLES[1]))
    PP.normalize_ooxml(os.path.join(golden, W.DELIVERABLES[1]), W.PEOPLE["demographer"], W.PEOPLE["director"],
                       "2026-10-13T14:05:00Z", "2026-10-16T19:42:18Z", company=W.SDC)
    json.dump(F, open(figures_path, "w", encoding="utf-8"), indent=1, sort_keys=False)
    open(figures_path, "a").write("\n")
    return X, F


if __name__ == "__main__":
    t = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(HERE), "target")
    main(t, os.path.join(os.path.dirname(t), "golden"), os.path.join(HERE, "figures.json"))
