"""The pack's documents: the program rules (PDF), the gas supply agreement excerpt (PDF), the data acquisition
system's export notes (txt), the generating unit register (xlsx) and the NOx allowance holdings (xlsx, a
distractor)."""
from reportlab import rl_config
rl_config.invariant = 1
from reportlab.lib import colors                                    # noqa: E402
from reportlab.lib.pagesizes import letter                          # noqa: E402
from reportlab.lib.styles import ParagraphStyle                     # noqa: E402
from reportlab.lib.units import inch                                # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate        # noqa: E402
from openpyxl import Workbook                                       # noqa: E402
from openpyxl.styles import Font                                    # noqa: E402

import world as W   # noqa: E402

SECTIONS = {"units": "2.1", "rate": "2.2", "designation": "2.3", "fuel": "3.1", "outages": "3.2"}
BODY = ParagraphStyle("b", fontName="Times-Roman", fontSize=10.5, leading=13.5, spaceAfter=6)
HEAD = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=11, leading=14, spaceBefore=8, spaceAfter=4)
TITLE = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=15, leading=19, spaceAfter=4)
SUB = ParagraphStyle("s", fontName="Helvetica", fontSize=9, leading=12, textColor=colors.HexColor("#444444"),
                     spaceAfter=10)


def _doc(path, title, org):
    d = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.95 * inch, rightMargin=0.95 * inch,
                          topMargin=0.9 * inch, bottomMargin=0.9 * inch, title=title, author=org,
                          creator=org, producer=org, subject=title)

    def frame(c, d_):
        c.saveState()
        c.setFont("Helvetica", 8)
        c.setFillColor(colors.HexColor("#555555"))
        c.drawString(0.95 * inch, 0.55 * inch, "%s  |  %s" % (org, title))
        c.drawRightString(letter[0] - 0.95 * inch, 0.55 * inch, "Page %d" % d_.page)
        c.restoreState()
    return d, frame


def write_guidelines(path):
    title = "SCR Upgrade Program 2027: Program Rules"
    d, frame = _doc(path, title, W.COOP)
    s = [Paragraph("Clean Air Commitment<br/>SCR Upgrade Program, 2027 Capital Budget: Program Rules", TITLE),
         Paragraph("%s, %s. Adopted by the Board of Directors %s." % (W.COOP, W.DEPT, W.GUIDELINE_DATE), SUB),
         Paragraph("1. Purpose", HEAD),
         Paragraph("Under its Clean Air Commitment the Cooperative holds each combined-cycle unit to an ozone-season "
                   "NOx rate target. The SCR Upgrade Program funds, in the 2027 capital budget, catalyst, reagent "
                   "injection and control upgrades on the units whose 2026 ozone-season rate did not meet it.", BODY),
         Paragraph("2. Designation", HEAD),
         Paragraph("2.1 <b>Units.</b> Every unit the generating unit register lists as a combined-cycle unit.", BODY),
         Paragraph("2.2 <b>Ozone-season NOx rate.</b> A unit's 2026 ozone-season NOx rate is the NOx the unit "
                   "emitted while operating from 1 May to 30 September 2026, in pounds, divided by its gross "
                   "generation in the same hours, in MWh. The hours are the operating hours of the data acquisition "
                   "system's hourly record, which is kept in standard time; hours of start-up and shutdown are "
                   "operating hours.", BODY),
         Paragraph("2.3 <b>Target.</b> A unit whose 2026 ozone-season NOx rate is 0.170 lb per MWh or more is "
                   "designated for the program.", BODY),
         Paragraph("3. The budget memo", HEAD),
         Paragraph("3.1 For every combined-cycle unit, the cost of the gas it burned in its 2026 ozone-season "
                   "operating hours, each hour's gas priced at the price on the supplier's invoice for the gas day the "
                   "hour falls in. A unit's gas is its heat input; one dekatherm is one MMBtu.", BODY),
         Paragraph("3.2 For every combined-cycle unit, its forced outage hours in the 2026 ozone season: the hours "
                   "from 1 May to 30 September 2026, local time, in which the outage event report shows the unit in a "
                   "full forced outage (event types U1, U2 and U3). Where the report holds a later report of the same outage, the "
                   "later report stands.", BODY),
         Paragraph("4. Timetable", HEAD),
         Paragraph("Environmental Services presents the designations to the Board with the 2027 capital budget on "
                   "%s." % W.BOARD_DATE, BODY)]
    d.build(s, onFirstPage=frame, onLaterPages=frame)


def write_terms(path):
    org = "Prairie Basin Gas Marketing LLC"
    title = "Firm Gas Supply Agreement: Excerpt"
    d, frame = _doc(path, title, org)
    s = [Paragraph("Firm Gas Supply Agreement between Prairie Basin Gas Marketing LLC and %s (excerpt)" % W.COOP,
                   TITLE),
         Paragraph("Effective 1 November 2024. Sections 1, 6 and 9.", SUB),
         Paragraph("Section 1. Definitions", HEAD),
         Paragraph("\"Day\" or \"Gas Day\" means a period of twenty-four consecutive hours beginning at 9:00 a.m. "
                   "Central Clock Time on a calendar day and ending at 9:00 a.m. Central Clock Time on the following "
                   "calendar day, and is identified by the calendar day on which it begins.", BODY),
         Paragraph("\"Central Clock Time\" means the prevailing time in the Central time zone, adjusted for daylight "
                   "saving time.", BODY),
         Paragraph("\"Transporter\" means the interstate pipeline that delivers Gas to a Delivery Point.", BODY),
         Paragraph("\"Delivery Point\" means the meter station serving each of Buyer's generating facilities listed "
                   "in Exhibit A.", BODY),
         Paragraph("Section 6. Price", HEAD),
         Paragraph("6.1 The price for Gas delivered on a Gas Day is the Daily Index for that Gas Day plus the charges of "
                   "the Transporter serving the Delivery Point, as set out for each Delivery Point in Exhibit A, in "
                   "dollars per dekatherm.", BODY),
         Paragraph("6.2 Quantities are the metered deliveries at each Delivery Point for the Gas Day, in dekatherms.",
                   BODY),
         Paragraph("Section 9. Billing", HEAD),
         Paragraph("9.1 Seller invoices monthly by Delivery Point.", BODY)]
    d.build(s, onFirstPage=frame, onLaterPages=frame)


def write_notes(path, n):
    lines = ["DAHS hourly export, calendar year 2026",
             "Exported 2026-10-19 from the plant data acquisition systems (DAHS) for every monitored unit.",
             "",
             "Rows: one per unit per clock hour of 2026, %d rows." % n,
             "OP_DATE, OP_HOUR: the hour beginning, in Central Standard Time all year (the DAHS clock does not",
             "  change for daylight saving time).",
             "OP_TIME: fraction of the hour the unit operated (0.00 when offline).",
             "GLOAD_MW: gross load, hourly average while operating. One operating hour at GLOAD_MW is GLOAD_MW MWh.",
             "HEAT_INPUT_MMBTU: from the fuel flow meter and the gas heating value.",
             "NOX_RATE_LB_MMBTU, NOX_MASS_LB: from the NOx and diluent analyzers; blank when the analyzers did not",
             "  produce a quality-assured hour.",
             "NOX_STATUS: OK quality-assured; OOC out of control (failed calibration or drift check); MAINT",
             "  maintenance; CALFAIL calibration in progress or failed; COMM communication loss; blank offline.",
             "This export is the raw DAHS record. It carries no substitute data; the quarterly reports to the",
             "  regulator are prepared from it separately.",
             "",
             "Contact: Environmental Services, DAHS administrator."]
    open(path, "w", encoding="utf-8", newline="\r\n").write("\n".join(lines) + "\n")


def write_units(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Units"
    ws.append(["Generating unit register, 2026"])
    ws.append(["%s, Generation Engineering. As of 30 June 2026." % W.COOP])
    ws.append([])
    ws.append(["Unit ID", "Facility", "Unit type", "Nameplate MW", "NOx control", "Commercial operation",
               "Dispatch name (before 2025)"])
    ws["A4"].font = Font(bold=True)
    old = {"CB-1": "CUTLER CC1", "CB-2": "CUTLER CC2", "AF-1": "ASHF GT1/ST1", "AF-2": "ASHF GT2/ST2",
           "DF-1": "DUNM 1", "DF-2": "DUNM 2", "DF-3": "DUNM 3", "SR-1": "SABLE CC1", "SR-2": "SABLE CC2",
           "AF-CT3": "ASHF CT3", "AF-CT4": "ASHF CT4", "SR-CT3": "SABLE CT3", "HC-4": "HOLT 4"}
    cod = {"CB-1": "2003-06-01", "CB-2": "2003-08-15", "AF-1": "2008-05-20", "AF-2": "2008-06-30", "DF-1": "2011-04-01",
           "DF-2": "2011-05-12", "DF-3": "2016-07-01", "SR-1": "2013-05-01", "SR-2": "2013-06-10",
           "AF-CT3": "1999-06-01", "AF-CT4": "1999-06-01", "SR-CT3": "2002-05-15", "HC-4": "1984-11-01"}
    for u, (plant, n, cap, *_r) in W.CC.items():
        ws.append([u, plant, "Combined cycle (1x1)", cap, "SCR (aqueous ammonia)", cod[u], old[u]])
    for u, (plant, kind, cap) in W.OTHER.items():
        ws.append([u, plant, kind.capitalize(), cap, "Water injection" if kind == "simple cycle" else "SNCR",
                   cod[u], old[u]])
    for col, w in zip("ABCDEFG", (9, 16, 22, 13, 16, 20, 26)):
        ws.column_dimensions[col].width = w
    wb.save(path)


def write_allowances(path, x):
    wb = Workbook()
    ws = wb.active
    ws.title = "Holdings"
    ws.append(["Cross-State Air Pollution Rule ozone-season NOx allowances: holdings by facility account"])
    ws.append(["Compliance year 2026. Allowances in tons. Reconciled to the allowance tracking system 1 October 2026."])
    ws.append([])
    ws.append(["Facility", "Allocated", "Purchased", "Transferred in", "Transferred out", "Held 1 Oct 2026"])
    rows = [("Cutler Bend Energy Center", 88, 0, 12, 0), ("Ashfork Generating Station", 104, 15, 0, 12),
            ("Dunmore Flats Power Plant", 131, 22, 0, 0), ("Sable Ridge Energy Center", 92, 0, 0, 0),
            ("Holt Creek Station", 512, 40, 0, 0)]
    for f, a, p, ti, to in rows:
        ws.append([f, a, p, ti, to, a + p + ti - to])
    wb.save(path)
