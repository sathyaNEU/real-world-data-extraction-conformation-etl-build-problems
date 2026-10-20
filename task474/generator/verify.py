"""The independent verifier: reads only target/ and golden/, recomputes the call and the terminal figures, and checks
the golden workbook and memo carry them."""
import os
import re
import sys

from openpyxl import load_workbook
from pypdf import PdfReader

import bundle as B
import methods as M

HERE = os.path.dirname(os.path.abspath(__file__))


def main(target, golden):
    fails = []

    def ok(label, cond):
        print("%s  verify: %s" % ("ok  " if cond else "FAIL", label))
        if not cond:
            fails.append(label)
    g = re.sub(r"\s+", " ", " ".join(p.extract_text() for p in PdfReader(B.find(target, "air_corridor")).pages))
    ok("pdf read back: rules 2.2, 2.3, 3.1", "times the class's NOx emission rate in the 2026 emission rate table" in g
       and "5.00 tonnes or more" in g and "under the motor carrier of its latest registration in force" in g)
    rate = B.nox(target)
    call = M.designated(rate)
    ok("the call: %s" % ", ".join(call), call == ["L03", "L05", "L14"])
    wb = load_workbook(B.find(golden, "air_corridor_links"), read_only=True)
    rows = {r[0]: r for r in wb["Links"].iter_rows(values_only=True) if r and str(r[0])[:1] == "L" and str(r[0])[1:].isdigit()}
    ok("the workbook lists the designated links: %s" % ", ".join(sorted(rows)), sorted(rows) == call)
    fq, fc = B.fleet_asks(target)
    rows = [r for r in wb["Drayage fleet"].iter_rows(min_row=4, values_only=True) if r and r[0] and r[1] is not None]
    want = [(a, b) for a, b in fq.values()] + [fc[c] for c in sorted(fc)]
    ok("fleet rows: %d" % len(rows), [(r[1], r[2]) for r in rows if isinstance(r[1], int)] == want)
    m = re.sub(r"\s+", " ", " ".join(p.extract_text() for p in PdfReader(B.find(golden, "air_corridor_designations")).pages))
    ok("the memo designates the three", "L03, L05 and L14" in m)
    print("verify: %d failures" % len(fails))
    return not fails


if __name__ == "__main__":
    task = os.path.dirname(HERE)
    sys.exit(0 if main(os.path.join(task, "target"), os.path.join(task, "golden")) else 1)
