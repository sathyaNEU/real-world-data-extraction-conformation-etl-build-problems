"""Independent verifier for task114. Reads only target/ and golden/ (and submission.md when present), shares no
code with the generator, recomputes every graded figure with plain Python, and checks the goldens against it.

    python3 verify_pack.py
"""
import csv
import datetime as dt
import re
import sys
import zipfile
from collections import defaultdict
from pathlib import Path
from zoneinfo import ZoneInfo

import pyarrow.parquet as pq
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
T = ROOT / "target"
GD = ROOT / "golden"
TOR, UTC = ZoneInfo("America/Toronto"), ZoneInfo("UTC")
checks = 0


def ok(cond, msg):
    global checks
    if not cond:
        sys.exit("FAIL: " + msg)
    checks += 1


def rows(name):
    with open(T / name, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def hm(s):
    return int(s[11:13]) * 60 + int(s[14:16])


def busiest(intervals):
    """Most intervals open at once; an interval ending at t does not meet one starting at t."""
    ev = []
    for s, e in intervals:
        ev.append((s, 1))
        ev.append((e, -1))
    ev.sort()
    c = best = 0
    for _, v in ev:
        c += v
        best = max(best, c)
    return best


def rule(peaks, k=12):
    """Fewest rooms short on no more than k days."""
    r = 0
    while sum(1 for v in peaks.values() if v > r) > k:
        r += 1
    return r


# ------------------------------------------------------------------ calendar, people, presence
work = [r["date"] for r in rows("firm_calendar_fy2026.csv") if r["day_type"] == "Working day"]
workset = set(work)
ok(len(work) == 246, "246 working days")
wb = load_workbook(T / "staff_directory_fy2026.xlsx", read_only=True)
ws = wb["Employees"]
hdr = [c.value for c in next(ws.iter_rows(max_row=1))]
emp_by_upn, emp_by_card, emp_by_badge = {}, {}, {}
for r in ws.iter_rows(min_row=2, values_only=True):
    d = dict(zip(hdr, r))
    emp_by_upn[d["UPN"].lower()] = d["Employee No"]
    emp_by_card[d["Adelaide Card"]] = d["Employee No"]
    emp_by_badge[d["Front St Badge"]] = d["Employee No"]
ok(len(emp_by_upn) == len(set(emp_by_upn.values())), "one UPN per employee")
present = defaultdict(set)
for r in rows("adelaide_access_events_fy2026.csv"):
    if r["result"] == "Granted" and r["card_number"] in emp_by_card:
        present[(emp_by_card[r["card_number"]], r["event_time"][:10])].add("A")
for r in rows("front_st_door_events_fy2026.csv"):
    if r["Badge"] in emp_by_badge:
        d = r["Date"]
        present[(emp_by_badge[r["Badge"]], "%s-%s-%s" % (d[6:10], d[3:5], d[0:2]))].add("F")
ok(all(len(v) == 1 for v in present.values()), "one office per person per day")

# ------------------------------------------------------------------ the main call
ev = pq.read_table(T / "calendar_events_fy2026.parquet").to_pylist()
keep = {}
for e in ev:
    if e["is_cancelled"] or e["is_all_day"] or e["categories"] in ("Client meeting", "Town hall"):
        continue
    day = e["start_local"][:10]
    if day not in workset:
        continue
    keep[e["event_id"]] = (day, hm(e["start_local"]), hm(e["end_local"]))
att = [r for r in rows("calendar_attendees_fy2026.csv") if r["event_id"] in keep]
NON_DECLINED = ("organizer", "accepted", "tentative", "none")


def replay(responses=NON_DECLINED, types=("required", "optional")):
    """Each meeting's invitees in each office that day, and the one-floor daily peaks."""
    n, na, nf = defaultdict(int), defaultdict(int), defaultdict(int)
    for r in att:
        if r["attendee_type"] not in types or r["response"] not in responses:
            continue
        emp = emp_by_upn.get(r["attendee_address"].lower())
        where = present.get((emp, keep[r["event_id"]][0])) if emp else None
        if where:
            n[r["event_id"]] += 1
            na[r["event_id"]] += "A" in where
            nf[r["event_id"]] += "F" in where
    byday = defaultdict(list)
    for eid, (day, s, e) in keep.items():
        if n[eid] >= 2:
            byday[day].append((s, e))
    return {d: busiest(byday.get(d, [])) for d in work}, n, na, nf


peaks, n, na, nf = replay()
R = rule(peaks)
for label, kw in (("accepted and organiser only", dict(responses=("organizer", "accepted"))),
                  ("accepted, tentative and organiser", dict(responses=("organizer", "accepted", "tentative"))),
                  ("tentative dropped, no response kept", dict(responses=("organizer", "accepted", "none"))),
                  ("declined invitees counted", dict(responses=NON_DECLINED + ("declined",))),
                  ("optional invitees ignored", dict(types=("required",)))):
    alt = rule(replay(**kw)[0])
    ok(abs(alt - R) / R >= 0.08, "reading '%s' (%d) lands at least 8 per cent from %d" % (label, alt, R))
short = lambda r: sum(1 for v in peaks.values() if v > r)
ok(short(R) <= 12 < short(R - 1), "the rule picks R")
split = sum(1 for eid in keep if n[eid] >= 2 and na[eid] < 2 and nf[eid] < 2)

# ------------------------------------------------------------------ the booking exports
rb = load_workbook(T / "room_register.xlsx", read_only=True)
rh = [c.value for c in next(rb["Rooms"].iter_rows(max_row=1))]
current = {}
mbox = {}
for r in rb["Rooms"].iter_rows(min_row=2, values_only=True):
    d = dict(zip(rh, r))
    current[d["Room Code"]] = d
    mbox[d["Mailbox"]] = d["Room Code"]
chh = [c.value for c in next(rb["Changes"].iter_rows(max_row=1))]
changes = defaultdict(list)
for r in rb["Changes"].iter_rows(min_row=2, values_only=True):
    d = dict(zip(chh, r))
    eff = d["Effective Date"].date() if hasattr(d["Effective Date"], "date") else d["Effective Date"]
    changes[d["Room Code"]].append((eff, d["Change"], d["Type Before"], d["Type After"]))
spells = defaultdict(list)
for code, d in current.items():
    start = dt.date(2019, 1, 1)
    for eff, what, before, after in sorted(changes[code]):
        if what != "Opened":
            spells[code].append((start, eff - dt.timedelta(days=1), before, d["Office"]))
        start = eff
    spells[code].append((start, None, d["Room Type"], d["Office"]))


def rtype(code, day):
    dd = dt.date.fromisoformat(day)
    for vf, vt, t, _ in spells[code]:
        if vf <= dd and (vt is None or dd <= vt):
            return t
    return None


# people of the firm on each item (rooms and guests aside); a series' UID carries the most any occurrence has
firm_on, firm_on_uid = defaultdict(int), defaultdict(int)
for r in rows("calendar_attendees_fy2026.csv"):
    if r["attendee_type"] in ("required", "optional") and r["attendee_address"].lower() in emp_by_upn:
        firm_on[r["event_id"]] += 1
uid_of_eid = {e["event_id"]: e["ical_uid"] for e in ev}
for eid, k in firm_on.items():
    firm_on_uid[uid_of_eid[eid]] = max(firm_on_uid[uid_of_eid[eid]], k)
client_eid = {e["event_id"] for e in ev if e["categories"] == "Client meeting"}
client_uid = {e["ical_uid"] for e in ev if e["categories"] == "Client meeting"}
focus_eid = {e["event_id"] for e in ev if e["subject"] == "Focus time"}
focus_uid = {e["ical_uid"] for e in ev if e["subject"] == "Focus time"}
halves = {c: d["Combinable With"].split(";") for c, d in current.items() if "+" in c and d["Combinable With"]}


def to_local(ts):
    return dt.datetime.strptime(ts, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=UTC).astimezone(TOR)


def adelaide(reading):
    """One hold per physical room a booking held: (day, start, end, one person, call, type, client).
    One person: one person of the firm on the booking's calendar item, a client meeting aside (an app or kiosk booking
    has no item, and its headcount says). A one-person booking is a call unless its calendar item is a focus block.
    correct: a booking ends where the platform gave the room back, and a combined booking holds both halves (its own
    row decides; the /1 /2 rows are its copies). legacy: release only on Released rows, combined rows dropped and the
    half rows kept."""
    out = []
    for r in rows("adelaide_room_bookings_fy2026.csv"):
        if r["status"] == "Cancelled":
            continue
        is_half, is_comb = "/" in r["booking_ref"], "+" in r["room_code"]
        if reading == "correct" and is_half or reading == "legacy" and is_comb:
            continue
        day = r["start"][:10]
        s, e = hm(r["start"]), hm(r["end"])
        if r["released_at"] and (reading == "correct" or r["status"] == "Released"):
            e = hm(r["released_at"])
        eid = r["outlook_event_id"]
        if reading == "legacy" or not eid:
            one = r["headcount"] == "1"
        else:
            one = firm_on.get(eid) == 1 and eid not in client_eid
        call = one and (reading == "legacy" or eid not in focus_eid)
        for code in (halves[r["room_code"]] if is_comb else [r["room_code"]]):
            out.append((day, s, e, one, call, rtype(code, day), r["outlook_event_id"] in client_eid))
    return out


def front(reading):
    """correct: a series row whose slot any exception took (whatever its status or room) is not a hold. legacy:
    cancelled and declined rows filtered first, and the series row dropped only where an exception took its slot in
    the same room."""
    out = []
    fb = rows("front_st_room_calendar_export.csv")
    live = ("Accepted", "Tentative")
    if reading == "correct":
        took = {(r["ical_uid"], r["original_start_utc"]) for r in fb if r["is_exception"] == "True"}
        gone = lambda r: (r["ical_uid"], r["start_utc"]) in took
    else:
        took = {(r["ical_uid"], r["original_start_utc"], r["room_mailbox"]) for r in fb
                if r["is_exception"] == "True" and r["response_status"] in live}
        gone = lambda r: (r["ical_uid"], r["start_utc"], r["room_mailbox"]) in took
    for r in fb:
        if r["response_status"] not in live or (r["is_exception"] == "False" and gone(r)):
            continue
        s, e = to_local(r["start_utc"]), to_local(r["end_utc"])
        day = s.date().isoformat()
        if reading == "legacy":
            one = r["attendee_count"] == "2"
        else:
            one = firm_on_uid.get(r["ical_uid"]) == 1 and r["ical_uid"] not in client_uid
        call = one and (reading == "legacy" or r["ical_uid"] not in focus_uid)
        out.append((day, s.hour * 60 + s.minute, e.hour * 60 + e.minute, one, call,
                    rtype(mbox[r["room_mailbox"]], day), r["ical_uid"] in client_uid))
    return out


holds = {"A": adelaide("correct"), "F": front("correct")}
legacy = {"A": adelaide("legacy"), "F": front("legacy")}

Q = [("Q1", "2025-10-01", "2025-12-31"), ("Q2", "2026-01-01", "2026-03-31"), ("Q3", "2026-04-01", "2026-06-30"),
     ("Q4", "2026-07-01", "2026-09-30")]
fig = {}
for o in ("A", "F"):
    mr = [h for h in holds[o] if h[5] == "meeting" and h[0] in workset]
    fig[o] = {}
    for q, a, b in Q:
        x = [h for h in mr if a <= h[0] <= b]
        days_q = [d for d in work if a <= d <= b]
        by = defaultdict(list)
        for h in x:
            by[h[0]].append((h[1], h[2]))
        fig[o][q] = dict(hours=round(sum(h[2] - h[1] for h in x) / 60.0, 2),
                         solo=round(sum(h[2] - h[1] for h in x if h[4]) / 60.0, 2),
                         peak=max(busiest(by.get(d, [])) for d in days_q))
    by = defaultdict(list)
    for h in mr:
        if not h[3] and not h[6]:
            by[h[0]].append((h[1], h[2]))
    fig[o]["own"] = rule({d: busiest(by.get(d, [])) for d in work})
    on = dt.date(2026, 9, 30)
    office = "Adelaide Street" if o == "A" else "Front Street"
    fig[o]["rooms"] = sum(1 for c, sp in spells.items() if "+" not in c and any(
        off == office and t == "meeting" and vf <= on and (vt is None or on <= vt) for vf, vt, t, off in sp))
both = defaultdict(list)
for o in ("A", "F"):
    for h in holds[o]:
        if h[5] == "meeting" and not h[3] and not h[6] and h[0] in workset:
            both[h[0]].append((h[1], h[2]))
natural = rule({d: busiest(both.get(d, [])) for d in work})
ok(natural < R, "the bookings on one timeline land below the replay (%d < %d)" % (natural, R))

# ------------------------------------------------------------------ the close-outs tie to the bookings
co = load_workbook(T / "room_utilisation_close_outs_fy2026.xlsx", read_only=True, data_only=True)
for q, a, b in Q:
    ws = co["%s FY26" % q]
    vals = {r[0]: r[1:3] for r in ws.iter_rows(min_row=5, max_row=8, values_only=True)}
    for i, o in enumerate(("A", "F")):
        for reading, H in (("correct", holds), ("legacy", legacy)):
            by = defaultdict(list)
            for h in H[o]:
                if h[5] in ("meeting", "client") and h[0] in workset and a <= h[0] <= b:
                    by[h[0]].append((h[1], h[2]))
            top = max(busiest(v) for v in by.values())
            first = min(d for d, v in by.items() if busiest(v) == top)
            pub = vals["Busiest moment: date"][i]
            pub = pub.date().isoformat() if hasattr(pub, "date") else str(pub)[:10]
            ok(vals["Busiest moment: rooms in use"][i] == top and pub == first,
               "close-out busiest moment %s %s on the %s reading" % (o, q, reading))
        ok(not any("hour" in str(k).lower() or "utilisation" in str(k).lower() for k in vals),
           "close-outs publish no hours %s %s" % (o, q))
        # rooms at quarter start: the physical count, never the count with the combined entries
        office = "Adelaide Street" if o == "A" else "Front Street"
        on = dt.date.fromisoformat(a)

        def rooms_on(kind, with_combined):
            return sum(1 for c, sp in spells.items() if (with_combined or "+" not in c) and any(
                off == office and t == kind and vf <= on and (vt is None or on <= vt) for vf, vt, t, off in sp))
        ok(vals["Bookable meeting rooms at quarter start"][i] == rooms_on("meeting", False),
           "close-out meeting rooms at quarter start are physical rooms %s %s" % (o, q))
        ok(vals["Bookable client rooms at quarter start"][i] == rooms_on("client", False),
           "close-out client rooms at quarter start %s %s" % (o, q))
        if o == "A":
            ok(rooms_on("meeting", True) != rooms_on("meeting", False),
               "counting combined entries misses the published count %s" % q)

# ------------------------------------------------------------------ the goldens
doc = zipfile.ZipFile(GD / "meeting_rooms_brief.docx").read("word/document.xml").decode("utf-8")
text = re.sub(r"<[^>]+>", "", doc)
ok("Meeting rooms for Level 21: %d" % R in text, "memo leads with the count")
ok("run short on %d working days" % short(R) in text and "run short on %d," % short(R - 1) in text,
   "memo carries days short at R and R-1")
ok("{:,}".format(split) in text, "memo carries the split-meeting count")
ok("gives %d rooms" % natural in text, "memo carries the bookings-only count")
for o in ("A", "F"):
    for q, _, _ in Q:
        for k in ("hours", "solo"):
            ok("{:,.2f}".format(fig[o][q][k]) in text, "memo carries %s %s %s" % (o, q, k))
    ok(str(fig[o]["own"]) in text and str(fig[o]["rooms"]) in text, "memo carries %s own count and rooms" % o)
above = [d for d, v in peaks.items() if v > R]
ok(all(dt.date.fromisoformat(d).weekday() in (1, 2) for d in above), "memo: every day above the count is a Tue or Wed")
ok(max(v for d, v in peaks.items() if dt.date.fromisoformat(d).weekday() in (0, 4)) < R / 2,
   "memo: Mondays and Fridays never come close")
ok("Tuesday or a Wednesday" in text, "memo carries the weekday sentence")
svg = (GD / "room_demand_fy2026.svg").read_text(encoding="utf-8")
ok("<svg" in svg and svg.count("<path") > 200, "chart is a vector drawing with its bars")
for f in [*T.iterdir(), *GD.iterdir()]:
    if f.suffix in (".md", ".csv", ".eml", ".svg"):
        ok("—" not in f.read_text(encoding="utf-8"), "no em dash in %s" % f.name)

# ------------------------------------------------------------------ submission.md: every number recomputed here
sub = ROOT / "submission.md"
if sub.exists():
    need_room = sum(1 for eid in keep if n[eid] >= 2)
    allowed = {str(R), str(R - 1), str(short(R)), str(short(R - 1)), str(len(work)), "{:,}".format(need_room),
               "{:,}".format(split), str(fig["A"]["rooms"] + fig["F"]["rooms"])}
    for o in ("A", "F"):
        allowed |= {str(fig[o]["own"]), str(fig[o]["rooms"])}
        for q, _, _ in Q:
            allowed |= {"{:,.2f}".format(fig[o][q]["hours"]), "{:,.2f}".format(fig[o][q]["solo"]),
                        str(fig[o][q]["peak"])}
    # structure, not findings: the task id, Level 21, the years and dates, headcount 1 and attendee_count 2
    structural = {"114", "21", "2025", "2026", "30", "1", "2"}
    text = sub.read_text(encoding="utf-8")
    body = re.sub(r"(?m)^(\s*|## )\d+\. ", "", text)  # list and block numbering are not figures
    for tok in re.findall(r"(?<![\w.])\d[\d,]*(?:\.\d+)?", body):
        tok = tok.rstrip(",")
        ok(tok in allowed or tok in structural, "submission figure %s recomputes" % tok)
    ok("—" not in text, "no em dash in submission.md")

# ------------------------------------------------------------------ house checks: registers, the tree, the archives
notes = (T / "extract_notes.md").read_text(encoding="utf-8")
for f in sorted(T.iterdir()):
    if f.name != "extract_notes.md":
        ok(f.name in notes, "extract notes register %s" % f.name)
ok(len(list(T.iterdir())) >= 10 and len({f.suffix for f in T.iterdir()}) >= 3, "input gates: 10+ files, 3+ formats")
ok(sorted(f.name for f in GD.iterdir()) == ["meeting_rooms_brief.docx", "room_demand_fy2026.svg"],
   "golden holds exactly the two deliverables the prompt names")
for name in ("submission.md", "prompt.md"):
    ok(len(list(ROOT.rglob(name))) == 1, "exactly one %s in the task tree" % name)
import hashlib
receipts = []
for zname, src in (("target.zip", T), ("golden.zip", GD)):
    zp = ROOT / zname
    if not zp.exists():
        continue
    with zipfile.ZipFile(zp) as z:
        names = z.namelist()
        ok(sorted(names) == sorted(f.name for f in src.iterdir()), "%s members match %s/ by name" % (zname, src.name))
        ok(not any("/" in n or n.startswith(".") or "__MACOSX" in n for n in names), "%s is flat and clean" % zname)
        for n in names:
            ok(z.read(n) == (src / n).read_bytes(), "%s member %s matches the tree" % (zname, n))
    receipts.append("%s  %d bytes  %d members  sha256 %s" % (zname, zp.stat().st_size, len(names),
                                                            hashlib.sha256(zp.read_bytes()).hexdigest()[:16]))

print("verify_pack: %d checks green" % checks)
for r in receipts:
    print(r)
print("committed count %d, days short %d at %d and %d at %d, natural stop %d, split meetings %d"
      % (R, short(R), R, short(R - 1), R - 1, natural, split))
for o in ("A", "F"):
    print(o, {q: fig[o][q] for q, _, _ in Q}, "own", fig[o]["own"], "rooms", fig[o]["rooms"])
