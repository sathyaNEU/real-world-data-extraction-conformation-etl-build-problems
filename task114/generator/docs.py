"""The pack's documents: the new office brief, the quarterly close-outs, the extract notes, the email thread."""
import datetime as dt
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib import colors

import golden as G

QNAME = {"Q1": "Q1 (October to December 2025)", "Q2": "Q2 (January to March 2026)",
         "Q3": "Q3 (April to June 2026)", "Q4": "Q4 (July to September 2026)"}


def brief(path):
    base = ParagraphStyle("b", fontName="Helvetica", fontSize=10, leading=13.5, spaceAfter=6)
    h = ParagraphStyle("h", parent=base, fontName="Helvetica-Bold", fontSize=11, spaceBefore=8, spaceAfter=4)
    t = ParagraphStyle("t", parent=base, fontName="Helvetica-Bold", fontSize=15, leading=19, spaceAfter=2)
    small = ParagraphStyle("s", parent=base, fontSize=8.5, leading=11, textColor=colors.HexColor("#555555"))
    doc = SimpleDocTemplate(str(path), pagesize=LETTER, leftMargin=2.3 * cm, rightMargin=2.3 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm, title="New office workplace brief, Rev C",
                            author="Workplace Planning")
    s = [Paragraph("Harrowfield Engineering &amp; Planning", small),
         Paragraph("New office workplace brief", t),
         Paragraph("Revision C, issued 2 October 2026 for Stage 3 design. From Workplace Planning to Corran Teague "
                   "Architects. Supersedes Revision B (21 August 2026).", small), Spacer(1, 8)]
    s.append(Paragraph("1. The move", h))
    s.append(Paragraph("Brennock Structures and Tallis Mobility Planning merged on 1 April 2025 and have traded as "
                       "Harrowfield since, from the Adelaide Street and Front Street offices. Both leases end in 2027. "
                       "From July 2027 the whole firm works from one floor, Level 21 of the new building on Queens "
                       "Quay East (4,700 m<super>2</super> net). Neither current office is kept.", base))
    s.append(Paragraph("2. Workstations", h))
    s.append(Paragraph("The workstation count was approved by the partners on 15 September 2026 from the attendance "
                       "study: 410 bookable workstations. It is not reopened by this revision.", base))
    s.append(Paragraph("3. Meeting rooms", h))
    s.append(Paragraph("Meeting rooms are sized on the reference year, 1 October 2025 to 30 September 2026. The new "
                       "office is given the fewest meeting rooms with which it would have run short of meeting rooms "
                       "on no more than twelve working days of the reference year. Working days are the days the firm "
                       "calendar marks as working days.", base))
    s.append(Paragraph("Room sizes come from the schedule of accommodation (2, 4, 6, 8, 12 and 20 seats). The mix is "
                       "agreed with the design team once the count is fixed.", base))
    s.append(Paragraph("A call taken by one person goes to the phone booths (28, in addition to the meeting rooms). "
                       "Client meetings are held in the client suite on the same floor (six rooms, sized separately). "
                       "The town hall space (160 seated) is outside the meeting-room count.", base))
    s.append(Paragraph("4. Systems", h))
    s.append(Paragraph("Room and workstation bookings in the new office run on the platform Adelaide Street uses today. "
                       "Front Street's room mailboxes are retired at the move.", base))
    s.append(Paragraph("5. Programme", h))
    tab = Table([["Stage 3 design freeze", "20 November 2026"], ["Fit-out", "January to June 2027"],
                 ["Move", "July 2027"]], colWidths=[6 * cm, 6 * cm])
    tab.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 9.5), ("LINEBELOW", (0, 0), (-1, -1), 0.4,
                                                                             colors.HexColor("#bbbbbb")),
                             ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
    s.append(tab)
    s.append(Spacer(1, 10))
    s.append(Paragraph("Not covered here: furniture, AV and signage, which follow in the Stage 3 report.", small))
    doc.build(s)


def close_outs(path, co):
    wb = Workbook()
    head = Font(bold=True)
    fill = PatternFill("solid", fgColor="DDE4EE")
    first = True
    for q in ("Q1", "Q2", "Q3", "Q4"):
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "%s FY26" % q
        ws["A1"] = "Room pressure close-out, %s" % QNAME[q]
        ws["A1"].font = Font(bold=True, size=12)
        ws["A2"] = "Facilities, prepared by K. Morgan. Distribution: partners, office managers."
        ws.append([])
        ws.append(["", "Adelaide Street", "Front Street"])
        for c in ws[4]:
            c.font = head
            c.fill = fill
        A, F = co["A"][q], co["F"][q]
        rows = [("Bookable meeting rooms at quarter start", A["meeting_rooms"], F["meeting_rooms"]),
                ("Bookable client rooms at quarter start", A["client_rooms"], F["client_rooms"]),
                ("Busiest moment: rooms in use", A["peak"], F["peak"]),
                ("Busiest moment: date", dt.date.fromisoformat(A["peak_date"]), dt.date.fromisoformat(F["peak_date"]))]
        for r in rows:
            ws.append(list(r))
        for row in ws.iter_rows(min_row=5, max_row=8):
            for c in row[1:]:
                if isinstance(c.value, float):
                    c.number_format = "#,##0.00" if "Hours" in str(row[0].value) else "0.0"
                if isinstance(c.value, dt.date):
                    c.number_format = "d mmm yyyy"
        ws.append([])
        ws.append(["Compiled from the two offices' booking systems at the quarter's close. Method on the last sheet."])
        ws.column_dimensions["A"].width = 44
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 18
    ws = wb.create_sheet("Method")
    lines = [
        "How the close-outs are compiled",
        "Scope: bookable meeting rooms and client rooms, each counted under the type it had on the day. Event "
        "spaces, project rooms and phone booths are not included.",
        "A room is held from the start of a booking to its end, or to the time the platform released it.",
        "A booking that holds more than one room counts each room it holds, once.",
        "Each occurrence of a recurring booking is counted at the time it took place.",
        "Bookings that were cancelled or declined are not counted.",
        "Times are Toronto local time. A room freed at a given minute is free for a booking that starts at that minute.",
        "Rooms in use at the busiest moment is the most rooms held at any one time on a working day in the quarter.",
        "Sources: Adelaide Street booking platform; Front Street room mailboxes; room register.",
    ]
    for i, ln in enumerate(lines):
        ws.cell(row=i + 1, column=1, value=ln)
        ws.cell(row=i + 1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 110
    wb.save(path)


NOTES = """# Workplace extracts for the new office brief

Pulled for Workplace Planning between 1 and 2 October 2026. Each extract covers 1 October 2025 to
30 September 2026 unless a line below says otherwise. Personal fields beyond work email and card numbers were
removed before the files were shared. Questions to Caleb Walker (IT) for the calendar pull, Kristie Morgan
(Facilities) for the room systems.

Also in this folder, as received: new_office_brief_stage2.pdf (Revision C), room_utilisation_close_outs_fy2026.xlsx
(Facilities' quarterly close-outs for the year) and RE_meeting_room_number.eml (the thread with the architect).

## calendar_events_fy2026.parquet
IT's pull for the Teams Rooms licence review (2 October 2026): every calendar item in the firm's tenant that
starts in the year, one row per occurrence.

- `event_id`: the item's id. `ical_uid`: shared by every occurrence of a series.
- `event_type`: singleInstance, occurrence (of a series) or exception (an occurrence changed from its series).
- `original_start_local`: for an exception, the start the series gave the occurrence.
- `start_local`, `end_local`: in `time_zone`. All-day items run midnight to midnight.
- `is_cancelled`: the organiser cancelled the item; it stays in the extract.
- `show_as`, `is_online_meeting`, `categories` as set in Outlook. `room_resources`: rooms on the item.

## calendar_attendees_fy2026.csv
One row per recipient of each item in the events file, rooms included.

- `attendee_type`: required, optional or resource (a room).
- `response`: organizer, accepted, tentative, declined or none (no response sent).
- Addresses outside harrowfield.ca are guests.

## adelaide_access_events_fy2026.csv
Adelaide Street card system, every reader, every event. `result` is Granted or Denied. Card numbers belong to
employees, contractors, cleaners and visitors; the staff directory carries employees' card numbers.

## front_st_door_events_fy2026.csv
Front Street door controller export. `Date` is day/month/year. `Badge` ids belong to employees (FS numbers)
and visitors (V numbers).

## staff_directory_fy2026.xlsx
HR extract of every employee on payroll at any time in the year, with leavers' end dates. Every employee
holds an Adelaide Street card and a Front Street badge.

## firm_calendar_fy2026.csv
Every date in the year with its day type: working day, statutory holiday, office closure or weekend.

## room_register.xlsx
Facilities' register of rooms in both offices. `Rooms` has one row per room as it stood when the register was
exported (1 October 2026). `Changes` logs every change to a room's type or availability since the offices were
fitted out, with the date it took effect. `Combinable With`: for a room that opens into its neighbour, the
neighbour; for the combined room, its two halves.

## adelaide_room_bookings_fy2026.csv
Adelaide Street booking platform export (1 October 2026).

- `status`: Confirmed, Released (nobody checked in and the platform released the room) or Cancelled.
- `checked_in_at`, `released_at`: Toronto local time. `released_at` is when the platform gave the room back.
- `headcount`: the number the booker entered, updated at check-in.
- `source`: Outlook, Mobile app or Kiosk. `outlook_event_id` is the calendar item for Outlook bookings.

## front_st_room_calendar_export.csv
Front Street room mailboxes, exported from Exchange on 1 October 2026. Times are UTC.

- `subject` shows the organiser's name (the mailboxes are set to hide meeting subjects).
- `response_status`: the room's response, Accepted, Tentative or Declined (the room was taken), or Cancelled.
- `is_exception`: True for an occurrence changed from its series. `original_start_utc`: its start in the series.
- `attendee_count`: the recipients on the request, the room included.
"""


def email(path, P):
    def msg(frm, to, cc, when, subject, body):
        return "From: %s\nTo: %s\n%sDate: %s\nSubject: %s\n\n%s\n" % (
            frm, to, ("Cc: %s\n" % cc) if cc else "", format_datetime(when), subject, body)

    tz = dt.timezone(dt.timedelta(hours=-4))
    lead = "Claudia Hernandez <claudia.hernandez@harrowfield.ca>"
    arch = "Mary Hurst <mhurst@corranteague.ca>"
    mp = "Louis Thompson <louis.thompson@harrowfield.ca>"
    co = "Debbie Jimenez <debbie.jimenez@harrowfield.ca>"
    parts = [
        msg(arch, lead, "", dt.datetime(2026, 10, 5, 9, 12, tzinfo=tz), "Meeting room number for the Stage 3 freeze",
            "Claudia,\n\nThanks for Rev C. The floor takes the 410 workstations with room to spare, and the core "
            "and the client suite are fixed. Everything else on the plan is waiting on the meeting room number, so "
            "we need it before the freeze on 20 November. The mix can follow a week or two later.\n\nMary\n\n"
            "Mary Hurst | Associate | Corran Teague Architects"),
        msg(mp, lead, co, dt.datetime(2026, 10, 5, 13, 40, tzinfo=tz), "RE: Meeting room number for the Stage 3 freeze",
            "Claudia,\n\nWalk either office mid-morning and you will find more empty rooms than full ones. I would "
            "rather not pay to build rooms we will not use, so please do not pad the number.\n\nLouis"),
        msg(co, lead, mp, dt.datetime(2026, 10, 5, 15, 5, tzinfo=tz), "RE: Meeting room number for the Stage 3 freeze",
            "For what it is worth, Front Street's rooms are at their busiest on Tuesday and Wednesday mornings. "
            "The Forum is only ever booked for the town halls.\n\nDebbie\nOffice "
            "Coordinator, Front Street"),
    ]
    body = "\n----- Original thread, newest last -----\n\n".join(parts)
    path.write_text(body, encoding="utf-8")


def write_docs(W, T):
    T = Path(T)
    brief(T / "new_office_brief_stage2.pdf")
    L = G.load(T)
    co = {"A": G.close_out_figures(L, G.adelaide_holds(L), "A"), "F": G.close_out_figures(L, G.front_holds(L), "F")}
    W.close_outs = co
    close_outs(T / "room_utilisation_close_outs_fy2026.xlsx", co)
    (T / "extract_notes.md").write_text(NOTES, encoding="utf-8")
    email(T / "RE_meeting_room_number.eml", W.people)
