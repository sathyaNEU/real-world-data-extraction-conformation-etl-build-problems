"""Render the two golden deliverables from the figures recomputed off target/ (golden_figures.json)."""
import datetime as dt
import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "golden"
SURFACE, INK, INK2, GRID = "#fcfcfb", "#0b0b0b", "#52514e", "#e4e3df"
BLUE, ORANGE = "#2a78d6", "#eb6834"
QN = {"Q1": "Oct to Dec 2025", "Q2": "Jan to Mar 2026", "Q3": "Apr to Jun 2026", "Q4": "Jul to Sep 2026"}


def d_long(s):
    d = dt.date.fromisoformat(s)
    return "%s %d %s %d" % (d.strftime("%A"), d.day, d.strftime("%B"), d.year)


def chart(g, path):
    R = g["main"]["rooms"]
    today = g["askA"]["today_rooms"]
    vals = sorted(g["daily_peaks"].values(), reverse=True)
    x = list(range(1, len(vals) + 1))
    cols = [ORANGE if v > R else BLUE for v in vals]
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "svg.fonttype": "path",
                         "svg.hashsalt": "task114"})
    f, ax = plt.subplots(figsize=(12.5, 6.4), dpi=100)
    f.patch.set_facecolor(SURFACE)
    ax.set_facecolor(SURFACE)
    ax.bar(x, vals, width=0.78, color=cols, linewidth=0, zorder=3)
    ax.axhline(R, color=INK, linewidth=1.6, zorder=4)
    ax.axhline(today, color=INK2, linewidth=1.3, linestyle=(0, (5, 3)), zorder=4)
    ax.text(len(vals) + 1, R + 0.6, "%d rooms, the committed count" % R, ha="right", va="bottom", fontsize=10,
            color=INK, zorder=5)
    ax.text(len(vals) + 1, today + 0.6, "%d meeting rooms today, Adelaide Street and Front Street together" % today,
            ha="right", va="bottom", fontsize=10, color=INK2, zorder=5)
    above = sum(1 for v in vals if v > R)
    ax.annotate("%d working days above %d rooms" % (above, R), xy=(above / 2 + 0.5, vals[0] + 0.3),
                xytext=(above + 14, vals[0] + 3.2), fontsize=10, color=INK, ha="left", va="center",
                arrowprops=dict(arrowstyle="-", color=INK2, linewidth=0.9))
    # the tail of the year is the Mondays and Fridays; label it only when the data says so exactly
    ranked = sorted(g["daily_peaks"].items(), key=lambda kv: (-kv[1], kv[0]))
    monfri = [dt.date.fromisoformat(d).weekday() in (0, 4) for d, _ in ranked]
    k = monfri.index(True) if True in monfri else len(monfri)
    if k < len(monfri) and all(monfri[k:]) and not any(monfri[:k]):
        top = max(v for _, v in ranked[k:])
        ax.annotate("", xy=(k + 0.6, top + 2.2), xytext=(len(vals) + 0.4, top + 2.2),
                    arrowprops=dict(arrowstyle="|-|", color=INK2, linewidth=0.9, mutation_scale=3))
        ax.text((k + len(vals)) / 2 + 0.5, top + 3.0, "Mondays and Fridays (%d days)" % (len(vals) - k),
                ha="center", va="bottom", fontsize=9.5, color=INK2)
    ax.set_xlim(0, len(vals) + 1.5)
    ax.set_ylim(0, today + 6)
    ax.set_xticks([1, 50, 100, 150, 200, len(vals)])
    ax.tick_params(axis="both", colors=INK2, length=0)
    ax.grid(axis="y", color=GRID, linewidth=0.8, zorder=0)
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.spines["bottom"].set_color("#c9c8c2")
    ax.set_xlabel("Working days of the reference year (1 Oct 2025 to 30 Sep 2026), busiest first", color=INK2)
    ax.set_ylabel("Meeting rooms in use at the day's busiest moment", color=INK2)
    f.text(0.012, 0.965, "Build %d meeting rooms: the new office would have run short on %d of %d working days"
           % (R, above, len(vals)), ha="left", va="top", fontsize=14, fontweight="bold", color=INK)
    f.text(0.012, 0.918, "Reference-year meetings replayed on the one floor, a room for every meeting with two or "
           "more of its people in that day", ha="left", va="top", fontsize=10.5, color=INK2)
    f.text(0.012, 0.018, "Source: calendar extract (IT, 2 October 2026), Adelaide Street card readers and Front Street "
           "door controller, 1 October 2025 to 30 September 2026. Workplace Planning.", ha="left", va="bottom",
           fontsize=8.5, color=INK2)
    f.subplots_adjust(left=0.065, right=0.985, top=0.86, bottom=0.13)
    f.savefig(path, format="svg", facecolor=SURFACE, metadata={"Date": None, "Creator": None})
    plt.close(f)
    import re
    svg = Path(path).read_text(encoding="utf-8")
    svg = svg.replace("matplotlib.axis_1", "days-axis").replace("matplotlib.axis_2", "rooms-axis")
    svg = re.sub(r"\s*<metadata>.*?</metadata>", "", svg, flags=re.S)
    Path(path).write_text(svg, encoding="utf-8")


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


def memo(g, path):
    R = g["main"]["rooms"]
    m = g["main"]
    A, F = g["askA"]["A"], g["askA"]["F"]
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2.0)

    def para(text, bold=False, size=None, space=6, color=None):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        p.paragraph_format.space_after = Pt(space)
        return p

    para("To: Mary Hurst, Corran Teague Architects\nFrom: Claudia Hernandez, Workplace Planning\n"
         "Date: 9 October 2026\nRe: Meeting rooms for Level 21 (Stage 3 freeze, 20 November)", size=9.5,
         color="52514E", space=10)
    para("Meeting rooms for Level 21: %d" % R, bold=True, size=15, space=4)
    para("Please plan Level 21 with %d meeting rooms. Had the new office been running through the reference year "
         "(1 October 2025 to 30 September 2026), %d rooms would have run short on %d working days, inside the "
         "brief's limit of twelve; %d rooms would have run short on %d, which is over it."
         % (R, R, m["days_short_at"], R - 1, m["days_short_one_fewer"]))
    para("How the number was set", bold=True, size=11.5, space=3)
    para("I replayed the reference year's calendar on the one floor. A meeting takes a room when two or more of its "
         "people who had not declined were in that day, in either building. That is the rule both offices' bookings "
         "follow all year: whenever two of a meeting's people were in the same building, there is a room booked "
         "there. Calls taken alone go to the booths, client meetings to the client suite and town halls to the "
         "town hall space, as the brief says. The count is the fewest rooms that would have run short on no more "
         "than twelve working days.")
    para("The room bookings alone point lower. Put both offices' bookings on one timeline and the same test gives "
         "%d rooms, but %s meetings had one of their people in each building that day, so neither office used a "
         "room for them. On Level 21 those people are on the same floor and the meetings need rooms. Every day "
         "above the count was a Tuesday or a Wednesday; Mondays and Fridays never came close."
         % (g["rungs"]["R2_bookings"], "{:,}".format(m["split_meetings"])))
    para("Against what we have today", bold=True, size=11.5, space=3)
    para("Meeting rooms only: client rooms, event spaces and project rooms are left out. Hours are the time rooms "
         "were held, from the start of a booking to its end or to its release; the calls column is the part of "
         "those hours one person booked for a call.", size=9.5, color="52514E", space=4)
    t = doc.add_table(rows=1, cols=7)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    hdr = ["Quarter", "Adelaide St hours", "of which calls", "Most rooms in use", "Front St hours",
           "of which calls", "Most rooms in use"]
    for c, h in zip(t.rows[0].cells, hdr):
        c.text = h
        c.paragraphs[0].runs[0].bold = True
        c.paragraphs[0].runs[0].font.size = Pt(9)
        shade(c, "DDE4EE")
    for q in ("Q1", "Q2", "Q3", "Q4"):
        row = t.add_row().cells
        vals = ["%s (%s)" % (q, QN[q]), "{:,.2f}".format(A[q]["hours"]), "{:,.2f}".format(A[q]["solo_hours"]),
                str(A[q]["peak"]), "{:,.2f}".format(F[q]["hours"]), "{:,.2f}".format(F[q]["solo_hours"]),
                str(F[q]["peak"])]
        for c, v in zip(row, vals):
            c.text = v
            c.paragraphs[0].runs[0].font.size = Pt(9)
            if v != vals[0]:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    para("Source: Adelaide Street booking platform and Front Street room mailboxes, both exported 1 October 2026; "
         "room types from Facilities' register.", size=8, color="52514E", space=8)
    t2 = doc.add_table(rows=1, cols=3)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t2.style = "Table Grid"
    for c, h in zip(t2.rows[0].cells, ["", "Adelaide Street", "Front Street"]):
        c.text = h
        if h:
            c.paragraphs[0].runs[0].bold = True
            c.paragraphs[0].runs[0].font.size = Pt(9)
        shade(c, "DDE4EE")
    for label, a, f in (("Bookable meeting rooms on 30 September 2026", A["rooms_30_sep"], F["rooms_30_sep"]),
                        ("Rooms its own bookings justify on the brief's rule", A["own_rule"], F["own_rule"])):
        row = t2.add_row().cells
        for c, v in zip(row, [label, str(a), str(f)]):
            c.text = v
            c.paragraphs[0].runs[0].font.size = Pt(9)
    para("Rooms as the register had them on 30 September 2026; AD-3.08 left booking and FR-2.05 joined it on 1 "
         "October. The Adelaide Street rooms that open in pairs (AD-3.06 with AD-3.07, AD-6.12 with AD-6.13) count "
         "as two rooms each, booked together or apart.", size=8, color="52514E", space=8)
    para("The two offices have %d meeting rooms between them today. Level 21 needs %d. The chart "
         "(room_demand_fy2026.svg) shows every working day of the reference year, busiest first, against both "
         "numbers." % (g["askA"]["today_rooms"], R), space=8)
    para("Claudia", space=0)
    fp = doc.sections[0].footer.paragraphs[0]
    fr = fp.add_run("Workplace Planning, Harrowfield Engineering & Planning        Page ")
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor.from_string("52514E")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "16")
    rpr.append(sz)
    r.append(rpr)
    tx = OxmlElement("w:t")
    tx.text = "1"
    r.append(tx)
    fld.append(r)
    fp._p.append(fld)
    doc.core_properties.author = "Claudia Hernandez"
    doc.core_properties.last_modified_by = "Claudia Hernandez"
    doc.core_properties.created = dt.datetime(2026, 10, 9, 10, 40)
    doc.core_properties.modified = dt.datetime(2026, 10, 9, 10, 40)
    doc.core_properties.revision = 1
    doc.core_properties.comments = ""
    doc.core_properties.title = "Meeting rooms for Level 21"
    doc.save(path)
    _fix_zip(path, dt.datetime(2026, 10, 9, 10, 40))


def _fix_zip(path, when):
    import io
    import zipfile
    src = zipfile.ZipFile(path)
    items = [(i.filename, src.read(i.filename)) for i in src.infolist()]
    src.close()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in items:
            zi = zipfile.ZipInfo(name, date_time=when.timetuple()[:6])
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, data)
    Path(path).write_bytes(buf.getvalue())


def main():
    g = json.loads((Path(__file__).parent / "golden_figures.json").read_text())
    OUT.mkdir(exist_ok=True)
    chart(g, OUT / "room_demand_fy2026.svg")
    memo(g, OUT / "meeting_rooms_brief.docx")


if __name__ == "__main__":
    main()
