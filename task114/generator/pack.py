"""Write the data files of target/ from the simulated world. Documents are written by docs.py."""
import base64
import csv
import datetime as dt
from collections import defaultdict
from pathlib import Path
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

from world import CLOSURE, DOMAIN, HOLIDAYS, OFFICES, D, T, employed

TOR = ZoneInfo("America/Toronto")
RESTRICTED = {"FR-3.34", "FR-3.35", "FR-3.36", "FR-4.35"}  # Front Street rooms whose mailbox needs a delegate
UTC = ZoneInfo("UTC")
ADL, FRT = 1, 2


def hhmm(day, minutes):
    return "%s %02d:%02d" % (day.isoformat(), minutes // 60, minutes % 60)


def utc(day, minutes):
    local = dt.datetime(day.year, day.month, day.day, minutes // 60, minutes % 60, tzinfo=TOR)
    return local.astimezone(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")


def mailbox(code, R):
    for r in R:
        if r["room_code"] == code:
            nm = r["room_name"].split(" (")[0].lower().replace(" ", "").replace("'", "")
            if code.startswith("AD"):
                return "rm-%s@%s" % (code.lower().replace("+", "-"), DOMAIN)
            return "fr.%s@%s" % (nm, DOMAIN)
    raise KeyError(code)


class Ids:
    def __init__(self, seed):
        self.rng = np.random.default_rng([seed, 99])
        self.used = set()

    def graph(self):
        while True:
            raw = bytes(self.rng.integers(0, 256, 27, dtype=np.uint8))
            s = "AAMkAD" + base64.urlsafe_b64encode(raw).decode().rstrip("=")
            if s not in self.used:
                self.used.add(s)
                return s

    def uid(self):
        while True:
            s = "040000008200E00074C5B7101A82E00800000000" + "".join(
                "%02X" % int(x) for x in self.rng.integers(0, 256, 16))
            if s not in self.used:
                self.used.add(s)
                return s


def texture_blocks(W, rng):
    """The calendar's texture in the order the pack draws it, the first draws on the pack's stream: out-of-office
    days, then each day's personal focus blocks with their start."""
    P, pres, out = W.people, W.pres, []
    for i, d in enumerate(W.work):
        for q in rng.choice(len(P), 9, replace=False):
            q = int(q)
            if employed(P[q], d) and pres[q, i] == 0:
                out.append(("x", i, q, d, 0))
        for q in rng.choice(len(P), 12, replace=False):
            q = int(q)
            if employed(P[q], d):
                out.append(("f", i, q, d, int(rng.choice([8 * 60, 13 * 60, 14 * 60]))))
    return out


def write_data(W, T_, seed=114):
    T_.mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng([seed, 7])
    ids = Ids(seed)
    P = W.people
    upn = {p["pid"]: p["upn"] for p in P}
    name = {p["pid"]: "%s %s" % (p["first"], p["last"]) for p in P}
    R = W.rooms
    rname = {r["room_code"]: r["room_name"] for r in R}

    # ------------------------------------------------ calendar extract (events and attendees)
    series_uid = {}
    ev_rows, at_rows = [], []
    rooms_of = defaultdict(list)  # meeting index -> room codes booked for it
    for b in W.bookings:
        if b["kind"] in ("meeting", "noshow"):
            rooms_of[("m", b["meeting"])].append(b["room"])
        elif b["kind"] == "client":
            rooms_of[("c", b["meeting"])].append(b["room"])
        elif b["kind"] == "townhall":
            rooms_of[("e", b["meeting"])].append(b["room"])
        elif b["kind"] == "solo" and b.get("outlook"):
            rooms_of[("s", b["meeting"])].append(b["room"])
        elif b["kind"] == "focus":
            rooms_of[b["meeting"]].append(b["room"])
    W.event_id = {}

    def add_event(key, day, start, dur, org, att, subject, category="", cancelled=False, series="",
                  pattern=None, ext=0, all_day=False, show_as="busy", online=True, guest_domain=None):
        if series:
            if series not in series_uid:
                series_uid[series] = ids.uid()
            uid = series_uid[series]
            typ = "exception" if pattern is not None else "occurrence"
        else:
            uid = ids.uid()
            typ = "singleInstance"
        eid = ids.graph()
        W.event_id[key] = eid
        rooms = sorted(set(rooms_of.get(key, [])))
        ev_rows.append(dict(
            event_id=eid, ical_uid=uid, event_type=typ, organizer=upn[org],
            subject=("Canceled: " + subject) if cancelled else subject,
            start_local=hhmm(day, start) if not all_day else day.isoformat() + " 00:00",
            end_local=hhmm(day, start + dur) if not all_day else (day + dt.timedelta(1)).isoformat() + " 00:00",
            time_zone="America/Toronto", is_all_day=all_day, is_cancelled=cancelled, show_as=show_as,
            is_online_meeting=online, categories=category,
            original_start_local=hhmm(day, pattern) if pattern is not None else "",
            room_resources=";".join(mailbox(c, R) for c in rooms)))
        for q, typ_, resp in att:
            at_rows.append((eid, upn[q], typ_, resp))
        for c in rooms:
            at_rows.append((eid, mailbox(c, R), "resource", "accepted"))
        for x in range(ext):
            dom = guest_domain or ["metrolinx-partner.ca", "cityworks.ca", "regionplan.ca", "hydrotek.ca"][x % 4]
            at_rows.append((eid, "guest%d.%s@%s" % (x + 1, uid[-6:].lower(), dom), "required",
                            "accepted" if x == 0 else "none"))

    for k, m in enumerate(W.meetings):
        pat = m["pattern_start"] if m["pattern_start"] is not None else (
            m["start"] if (m["cancelled"] and m["series"]) else None)
        add_event(("m", k), m["day"], m["start"], m["dur"], m["org"], m["att"], m["subject"],
                  cancelled=m["cancelled"], series=m["series"], pattern=pat)
    for k, c in enumerate(W.client):
        add_event(("c", k), c["day"], c["start"], c["dur"], c["org"], c["att"], c["subject"], category="Client meeting",
                  ext=c["ext"], online=False)
    for k, e in enumerate(W.events):
        add_event(("e", k), e["day"], e["start"], e["dur"], e["org"], e["att"], "Town hall", category="Town hall")
    for k, s in enumerate(W.solo):
        if s["outlook"]:
            # a call with someone outside the firm on the line carries them on the invite
            add_event(("s", k), s["day"], s["start"], s["dur"], s["org"], [(s["org"], "required", "organizer")],
                      s.get("subject", "Call"), online=True, series=s["series"], pattern=s["pattern_start"],
                      ext=s.get("guests", 0), guest_domain=s.get("domain"))
    for c in sorted(W.standing_cancelled, key=lambda c: (c["day"], c["s"], c["sid"])):
        add_event(("sc", c["sid"], c["day"]), c["day"], c["s"], c["e"] - c["s"], c["org"],
                  [(c["org"], "required", "organizer")], "Call", cancelled=True, online=True, series=c["sid"],
                  pattern=c["s"])
    # texture: out-of-office all-day entries and personal focus blocks, which hold no meeting
    pres = W.pres
    blocks = texture_blocks(W, rng)
    assert [b for b in blocks if b[0] == "f"] == W.focus_blocks, "the planner saw the focus blocks the pack writes"
    for kind, i, q, d, st in blocks:
        if kind == "x":
            add_event(("x", i, q), d, 0, 1440, q, [(q, "required", "organizer")], "Out of office", all_day=True,
                      show_as="oof", online=False)
        else:
            add_event(("f", i, q), d, st, 120, q, [(q, "required", "organizer")], "Focus time", online=False)
    ev = pd.DataFrame(ev_rows).sort_values(["start_local", "organizer", "event_id"]).reset_index(drop=True)
    ev.to_parquet(T_ / "calendar_events_fy2026.parquet", index=False, engine="pyarrow", compression="snappy")
    at = pd.DataFrame(at_rows, columns=["event_id", "attendee_address", "attendee_type", "response"])
    order = {e: i for i, e in enumerate(ev["event_id"])}
    at["o"] = at["event_id"].map(order)
    at = at.sort_values(["o", "attendee_type", "attendee_address"], kind="mergesort").drop(columns="o")
    at.to_csv(T_ / "calendar_attendees_fy2026.csv", index=False, lineterminator="\n")

    # ------------------------------------------------ access events
    a_rows, f_rows = [], []
    a_readers = ["ADL L1 Turnstile 1", "ADL L1 Turnstile 2", "ADL L1 Turnstile 3", "ADL P1 Parking Door"]
    a_floor = ["ADL L3 East Door", "ADL L4 East Door", "ADL L5 West Door", "ADL L6 West Door", "ADL L7 Reception"]
    f_doors = ["FRT Lobby Gate A", "FRT Lobby Gate B", "FRT Loading Dock"]
    f_floor = ["FRT L2 Stair", "FRT L3 Suite Door", "FRT L4 Suite Door"]
    for i, d in enumerate(W.work):
        for p in P:
            o = pres[p["pid"], i]
            if o == 0:
                continue
            arr = int(rng.integers(7 * 60 + 15, 10 * 60 + 40))
            sec = int(rng.integers(0, 60))
            if o == ADL:
                card = p["adelaide_card"]
                rd = a_readers[int(rng.choice(4, p=[0.34, 0.33, 0.28, 0.05]))]
                if rng.random() < 0.02:
                    a_rows.append((d, arr * 60 + sec, rd, card, "Denied"))
                    sec = min(59, sec + 4)
                a_rows.append((d, arr * 60 + sec, rd, card, "Granted"))
                for _ in range(int(rng.integers(0, 3))):
                    t = int(rng.integers(arr + 5, 17 * 60))
                    a_rows.append((d, t * 60 + int(rng.integers(0, 60)), a_floor[int(rng.integers(5))], card, "Granted"))
                if rng.random() < 0.3:
                    t = int(rng.integers(12 * 60, 14 * 60))
                    a_rows.append((d, t * 60 + int(rng.integers(0, 60)), a_readers[int(rng.integers(3))], card,
                                   "Granted"))
            else:
                badge = p["front_badge"]
                door = f_doors[int(rng.choice(3, p=[0.55, 0.42, 0.03]))]
                f_rows.append((d, arr * 60 + sec, door, badge, "IN"))
                for _ in range(int(rng.integers(0, 2))):
                    t = int(rng.integers(arr + 5, 17 * 60))
                    f_rows.append((d, t * 60 + int(rng.integers(0, 60)), f_floor[int(rng.integers(3))], badge, "IN"))
                lv = int(rng.integers(15 * 60 + 30, 19 * 60))
                f_rows.append((d, lv * 60 + int(rng.integers(0, 60)), f_doors[int(rng.integers(2))], badge, "OUT"))
        # contractors, cleaners and visitors, whose cards are not on the staff directory
        for _ in range(int(rng.integers(8, 18))):
            t = int(rng.integers(6 * 60, 18 * 60))
            a_rows.append((d, t * 60 + int(rng.integers(0, 60)), a_readers[int(rng.integers(4))],
                           "%08d" % int(rng.integers(18000000, 18999999)), "Granted"))
        for _ in range(int(rng.integers(6, 14))):
            t = int(rng.integers(6 * 60, 18 * 60))
            f_rows.append((d, t * 60 + int(rng.integers(0, 60)), f_doors[int(rng.integers(3))],
                           "V%05d" % int(rng.integers(1, 99999)), "IN"))
    a_rows.sort(key=lambda r: (r[0], r[1], r[2], r[3]))
    with open(T_ / "adelaide_access_events_fy2026.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["event_time", "reader_name", "card_number", "result"])
        for d, s, rd, card, res in a_rows:
            w.writerow(["%s %02d:%02d:%02d" % (d.isoformat(), s // 3600, s // 60 % 60, s % 60), rd, card, res])
    f_rows.sort(key=lambda r: (r[0], r[1], r[2], r[3]))
    with open(T_ / "front_st_door_events_fy2026.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["Date", "Time", "Door", "Badge", "Direction"])
        for d, s, door, badge, dirn in f_rows:
            w.writerow([d.strftime("%d/%m/%Y"), "%02d:%02d:%02d" % (s // 3600, s // 60 % 60, s % 60), door, badge, dirn])

    # ------------------------------------------------ staff directory
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    hdr = ["Employee No", "Display Name", "UPN", "Legacy Firm", "Home Office", "Practice", "Team", "Grade",
           "Manager Employee No", "Start Date", "End Date", "Adelaide Card", "Front St Badge", "Contract"]
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True)
    emp_no = {p["pid"]: p["employee_no"] for p in P}
    for p in sorted(P, key=lambda p: p["employee_no"]):
        ws.append([p["employee_no"], "%s, %s" % (p["last"], p["first"]), p["upn"],
                   "Brennock Structures" if p["office"] == ADL else "Tallis Mobility Planning",
                   OFFICES[p["office"]], p["practice"], p["team"], p["grade"],
                   emp_no.get(W.manager.get(p["pid"]), ""), p["start"], p["end"], p["adelaide_card"],
                   p["front_badge"], "Part time (4 days)" if p["part_time"] else "Full time"])
    for row in ws.iter_rows(min_row=2):
        for c in row[9:11]:
            c.number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    wb.save(T_ / "staff_directory_fy2026.xlsx")

    # ------------------------------------------------ firm calendar
    with open(T_ / "firm_calendar_fy2026.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["date", "weekday", "day_type", "note"])
        for d in W.days:
            if d.weekday() >= 5:
                dtp, note = "Weekend", ""
            elif d in HOLIDAYS:
                dtp, note = "Statutory holiday", HOLIDAYS[d]
            elif d in CLOSURE:
                dtp, note = "Office closure", "Offices closed between Christmas and New Year"
            else:
                dtp, note = "Working day", ""
            w.writerow([d.isoformat(), d.strftime("%a"), dtp, note])

    # ------------------------------------------------ room register: rooms as they stand, and the changes log
    wb = Workbook()
    ws = wb.active
    ws.title = "Rooms"
    ws.append(["Room Code", "Mailbox", "Office", "Floor", "Room Name", "Seats", "Room Type", "Combinable With"])
    for c in ws[1]:
        c.font = Font(bold=True)
    last = {}
    for r in sorted(R, key=lambda r: (r["office"], r["room_code"], r["valid_from"])):
        last[r["room_code"]] = r
    for code in sorted(last, key=lambda c: (last[c]["office"], c)):
        r = last[code]
        ws.append([r["room_code"], mailbox(r["room_code"], R), r["office"], r["floor"], r["room_name"], r["seats"],
                   r["room_type"], r["combinable_with"]])
    ws.freeze_panes = "A2"
    ch = wb.create_sheet("Changes")
    ch.append(["Effective Date", "Room Code", "Change", "Type Before", "Type After", "Raised By"])
    for c in ch[1]:
        c.font = Font(bold=True)
    spells = defaultdict(list)
    for r in R:
        spells[r["room_code"]].append(r)
    events = []
    for code, sp in spells.items():
        sp = sorted(sp, key=lambda r: r["valid_from"])
        if sp[0]["valid_from"] > D(2019, 1, 1):
            events.append((sp[0]["valid_from"], code, "Opened", "", sp[0]["room_type"], "K. Morgan"))
        for a, b in zip(sp, sp[1:]):
            what = "Handed to the client suite" if b["room_type"] == "client" else (
                "Released from the client suite" if a["room_type"] == "client" else "Taken out of booking")
            events.append((b["valid_from"], code, what, a["room_type"], b["room_type"], "K. Morgan"))
    for e in sorted(events):
        ch.append(list(e))
    for row in ch.iter_rows(min_row=2):
        row[0].number_format = "yyyy-mm-dd"
    ch.column_dimensions["A"].width = 14
    ch.column_dimensions["C"].width = 30
    wb.save(T_ / "room_register.xlsx")

    # ------------------------------------------------ booking exports
    held = W.rooms_held.held
    by_day = defaultdict(dict)
    for (c, dd), iv in held.items():
        by_day[dd][c] = iv
    a_out, f_out = [], []
    ghost_of = defaultdict(list)
    for gh in W.ghosts:
        if gh["why"] in ("moved", "cancelled"):
            ghost_of[id(gh["src"])].append(gh)
    ref = 104000
    for b in sorted(W.bookings, key=lambda b: (b["day"], b["s"], b["room"], b["kind"], b["meeting"])):
        day = b["day"]
        if b["office"] == ADL:
            ref += int(rng.integers(1, 4))
            key = {"meeting": "m", "noshow": "m", "cancelled": "m", "client": "c", "townhall": "e",
                   "solo": "s"}.get(b["kind"])
            eid = W.event_id.get(b["meeting"], "") if b["kind"] == "focus" else (
                W.event_id.get((key, b["meeting"]), "") if key else "")
            src = "Outlook" if (b["kind"] != "solo" or b.get("outlook")) else (
                "Mobile app" if rng.random() < 0.8 else "Kiosk")
            created = T(day.year, day.month, day.day) - dt.timedelta(days=int(rng.integers(0, 9)),
                                                                     minutes=int(rng.integers(30, 600)))
            if b["kind"] == "solo" and not b.get("outlook"):
                created = T(day.year, day.month, day.day) + dt.timedelta(minutes=b["s"] - int(rng.integers(5, 60)))
            status = {"noshow": "Released", "cancelled": "Cancelled"}.get(b["kind"], "Confirmed")
            ci = ""
            if status == "Confirmed":
                ci = (T(day.year, day.month, day.day) + dt.timedelta(minutes=b["s"] - 3 + int(rng.integers(0, 12)))
                      ).strftime("%Y-%m-%d %H:%M")
            rel = hhmm(day, b["released"]) if status == "Released" else ""
            if b.get("gave_back") is not None:
                rel = hhmm(day, b["gave_back"])
            hc = b["n"] if b["kind"] != "townhall" else 0
            row = dict(booking_ref="ADB-%06d" % ref, room_code=b["room"], room_name=rname[b["room"]],
                       start=hhmm(day, b["s"]), end=hhmm(day, b["e"]), booked_for=upn[b["org"]], headcount=hc,
                       status=status, checked_in_at=ci, released_at=rel, source=src, outlook_event_id=eid,
                       created_at=created.strftime("%Y-%m-%d %H:%M"))
            a_out.append(row)
            for j, h in enumerate([] if b.get("one_row") else b["halves"]):
                r2 = dict(row)
                r2.update(booking_ref="ADB-%06d/%d" % (ref, j + 1), room_code=h, room_name=rname[h])
                a_out.append(r2)
        else:
            key = {"meeting": "m", "cancelled": "m", "client": "c", "townhall": "e", "solo": "s", "focus": "f"}[b["kind"]]
            eid_key = b["meeting"] if key == "f" else (key, b["meeting"])
            uid = None
            if key == "m":
                m = W.meetings[b["meeting"]]
                uid = series_uid.get(m["series"]) if m["series"] else None
            elif key == "s" and W.solo[b["meeting"]]["series"]:
                uid = series_uid.get(W.solo[b["meeting"]]["series"])
            if uid is None:
                uid = _uid_of(W, eid_key, ev)
            n_recip = 2 if key == "f" else _recipients(W, key, b["meeting"], rooms_of)
            base = dict(room_mailbox=mailbox(b["room"], R), subject=name[b["org"]], organizer=upn[b["org"]],
                        start_utc=utc(day, b["s"]), end_utc=utc(day, b["e"]),
                        response_status="Cancelled" if b["kind"] == "cancelled" else "Accepted",
                        is_recurring=bool(b["series"]),
                        is_exception=False, original_start_utc="", attendee_count=n_recip, ical_uid=uid)
            if b["kind"] == "cancelled" and b["series"]:
                # a cancelled occurrence of a series is an exception; the request named the series' room
                base.update(is_exception=True, original_start_utc=utc(day, b["s"]), attendee_count=n_recip + 1)
                f_out.append(base)
                for gh in ghost_of.get(id(b), []):
                    g = dict(base)
                    g.update(room_mailbox=mailbox(gh["room"], R), response_status="Accepted", is_exception=False,
                             original_start_utc="")
                    f_out.append(g)
                continue
            if b["room"] in RESTRICTED and b["kind"] in ("meeting", "solo") and rng.random() < 0.35:
                base["response_status"] = "Tentative"
            if b["kind"] in ("meeting", "solo") and b["pattern"] is not None:
                ps = b["pattern"]
                base["is_exception"] = True
                base["original_start_utc"] = utc(day, ps)
                f_out.append(base)
                # the series row for the slot the occurrence left stays in the series' room
                for gh in ghost_of.get(id(b), []):
                    g = dict(base)
                    g.update(room_mailbox=mailbox(gh["room"], R), start_utc=utc(day, gh["s"]),
                             end_utc=utc(day, gh["e"]), is_exception=False, original_start_utc="")
                    f_out.append(g)
                    b["ghost"] = True
                continue
            if b["room"] in RESTRICTED and b["kind"] in ("meeting", "solo") and rng.random() < 0.35:
                base["response_status"] = "Tentative"  # approval-only rooms: the delegate never actioned it
            f_out.append(base)
            # an occasional first request to a room that was already taken comes back declined
            if b["kind"] == "meeting" and rng.random() < 0.04:
                busy = [c for c, iv in by_day[day].items() if c.startswith("FR") and c != b["room"]
                        and any(not (b["e"] <= s or e <= b["s"]) for s, e in iv)]
                if busy:
                    c = sorted(busy)[int(rng.integers(len(busy)))]
                    dcl = dict(base)
                    dcl.update(room_mailbox=mailbox(c, R), response_status="Declined")
                    f_out.append(dcl)
    # standing calls cancelled on a day their owner was out: the exception, and the series row where it stayed
    for c in W.standing_cancelled:
        day = c["day"]
        base = dict(room_mailbox=mailbox(c["room"], R), subject=name[c["org"]], organizer=upn[c["org"]],
                    start_utc=utc(day, c["s"]), end_utc=utc(day, c["e"]), response_status="Cancelled",
                    is_recurring=True, is_exception=True, original_start_utc=utc(day, c["s"]), attendee_count=2,
                    ical_uid=series_uid[c["sid"]])
        f_out.append(base)
        if c["ghost"]:
            g = dict(base)
            g.update(response_status="Accepted", is_exception=False, original_start_utc="")
            f_out.append(g)
    a = pd.DataFrame(a_out)
    a.to_csv(T_ / "adelaide_room_bookings_fy2026.csv", index=False, lineterminator="\n")
    f = pd.DataFrame(f_out).sort_values(["start_utc", "room_mailbox", "response_status"], kind="mergesort")
    f.to_csv(T_ / "front_st_room_calendar_export.csv", index=False, lineterminator="\n")
    return dict(events=len(ev), attendees=len(at), a_access=len(a_rows), f_access=len(f_rows), a_book=len(a),
                f_book=len(f))


def _uid_of(W, key, ev):
    eid = W.event_id.get(key)
    if eid is None:
        return ""
    if not hasattr(W, "_uid_by_eid"):
        W._uid_by_eid = dict(zip(ev["event_id"], ev["ical_uid"]))
    return W._uid_by_eid[eid]


def _recipients(W, key, k, rooms_of):
    if key == "m":
        att = W.meetings[k]["att"]
    elif key == "c":
        c = W.client[k]
        return len(c["att"]) + c["ext"] + len(set(rooms_of.get(("c", k), [])))
    elif key == "e":
        att = W.events[k]["att"]
    else:
        return 1 + W.solo[k].get("guests", 0) + len(set(rooms_of.get(("s", k), [])))
    return len(att) + len(set(rooms_of.get((key, k), [])))
