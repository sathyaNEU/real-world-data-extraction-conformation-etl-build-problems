"""Build the pack and the goldens, write metadata.json and figures.json, run every check and the independent
verifier, build a second time into a scratch folder and require byte-identical output, then cut target.zip.

    python3 task254/generator/make.py [--once]
"""
import filecmp
import json
import os
import shutil
import sys
import tempfile
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import analysis as AN
import asks as A
import build as B
import checks as C
import core
import docs as D
import golden as GD
import pack as P
import verify as V
import world as WD

TASK = os.path.dirname(HERE)
TARGET = os.path.join(TASK, "target")
GOLDEN = os.path.join(TASK, "golden")
STAMP = 1792108800          # 16 October 2026

SOURCES = {
    "tmc": "Probe vendor's TMC identification table, network releases 2025.2 and 2026.1",
    "readme": "Probe vendor's readme for the TMC travel-time export (delivered 5 October 2026)",
    "route": "Transportation Planning's route file for the morning staff shuttle, Cedar Grove to Brook Street",
    "deadhead": "Shuttle operator's deadhead route, Brook Street stop to its Fern Valley yard (October 2025)",
    "standard": "Staff shuttle service standard, version 2 (June 2026)",
    "budget": "Finance note on the 2027 budget basis of the staff shuttle (12 October 2026)",
    "notice": "Parking Services notice of the Chestnut Street garage closure (15 September 2026)",
    "charges": "Shuttle operator's 2027 schedule of charges (18 September 2026)",
    "holidays": "HR hospital holiday calendars, October 2025 to September 2026 and calendar 2027",
    "hr": "HR staff extract of 2 October 2026: identifiers, department, cost centre, scheduled start, FTE",
    "reassign": "Parking Services reassignment workbook: permit roll, HR-granted spaces and opt-outs",
    "survey": "2026 staff commute survey responses",
    "transit": "Regional transit authority's Route 40X express timetable (effective 17 August 2026)",
}


def source_of(k):
    if k.startswith("probe_"):
        return "Probe vendor's TMC travel-time export for %s: every segment in the export, every day, 05:30 to 09:25" \
               % k[6:]
    return SOURCES[k]


def write_all(target, golden):
    if os.path.isdir(target):
        shutil.rmtree(target)
    W, S = B.build()
    P.write(W, S, target)
    D.standard_docx(os.path.join(target, P.FILES["standard"]))
    D.budget_docx(os.path.join(target, P.FILES["budget"]))
    D.notice_pdf(os.path.join(target, P.FILES["notice"]))
    D.transit_pdf(os.path.join(target, P.FILES["transit"]))
    D.readme_txt(os.path.join(target, P.FILES["readme"]), P.FILES)
    F = AN.main_figures(W)
    G = AN.ask_figures(F, S)
    GD.write(golden, F, G)
    for f in os.listdir(target):
        os.utime(os.path.join(target, f), (STAMP, STAMP))
    return W, S, F, G


def rows_of(path):
    if path.endswith(".csv"):
        with open(path) as f:
            return sum(1 for _ in f) - 1
    return None


def metadata():
    files = []
    for k, f in sorted(P.FILES.items(), key=lambda x: x[1]):
        pth = os.path.join(TARGET, f)
        files.append({"path": f, "format": f.rsplit(".", 1)[1], "bytes": os.path.getsize(pth), "rows": rows_of(pth),
                      "source": "Constructed: " + source_of(k), "date": "2026-10-16",
                      "license": "CC BY 4.0 (original work, constructed for this task)"})
    big = max((x for x in files if x["rows"]), key=lambda x: x["rows"])
    meta = {
        "task": "task254",
        "title": "The latest 2027 morning departure of a hospital's staff shuttle from a remote park-and-ride lot "
                 "whose trips would have reached the campus stop by 8:00 a.m. on 90 per cent of working weekdays",
        "domain": "Urban & Regional Science",
        "subdomain": "network analysis (corridor travel-time reliability)",
        "objective": "Descriptive & Distribution Analysis",
        "prompt_shape": "04 (setting one dial): the latest departure meeting the on-time share, the working weekdays "
                        "at it and the next departure, and at each departure from 6:50 to 7:20 the late staff "
                        "arrivals and the 2027 cost of holding it, with the riders, vehicles and seats it carries",
        "as_of": "2026-10-16",
        "deliverables": [GD.XLSX, GD.HTML],
        "distractor_files": [P.FILES[k] for k in P.DISTRACTORS],
        "input_gates": {"files": len(files), "formats": sorted({x["format"] for x in files}),
                        "largest_file": big["path"], "largest_file_rows": big["rows"]},
        "sources": [{"files": "all files under target/",
                     "origin": "constructed for this task by task254/generator/make.py (seeded, deterministic) "
                               "around a fictional hospital; I-65, I-264 and the Louisville streets are named as "
                               "real roads, the probe vendor, the shuttle operator, the hospital, its staff and "
                               "every travel time are invented, and no real probe data is reproduced.",
                     "date": "2026-10-16", "license": "CC BY 4.0 (original work)"}],
        "created": "2026-10-12",
        "generator": "task254/generator/make.py",
        "files": files,
    }
    with open(os.path.join(TASK, "metadata.json"), "w") as f:
        json.dump(meta, f, indent=1)
        f.write("\n")


def hm(m):
    return "%d:%02d" % (m // 60, m % 60)


def ampm(m):
    return "%d:%02d a.m." % (m // 60, m % 60)


def figures_json(W, S, F, G):
    call = F["call"]
    gold = G["gold"]
    deps = core.DEPARTURES
    late_days = {d: F["N"] - F["golden_counts"][d] for d in deps}
    out = {
        "call": ampm(call),
        "call_next": ampm(call + 5),
        "on_time_at_call": F["golden_counts"][call],
        "on_time_at_next": F["golden_counts"][call + 5],
        "working_weekdays": F["N"],
        "bar_weekdays": F["need"],
        "bar_pct": 90,
        "bar_exact": "%.1f" % F["bar"],
        "arrive_by": "8:00 a.m.",
        "window": ["1 October 2025", "30 September 2026"],
        "legs_minutes": [core.LOT_LEG, core.STOP_LEG],
        "route_segments": len(core.ROUTE),
        "holidays_2025_26": len(core.HOLIDAYS),
        "riders": gold["riders"],
        "vehicles": gold["vehicles"],
        "seats_empty": gold["seats_empty"],
        "seats_per_vehicle": A.SEATS,
        "riders_both_sheets": G["both_sheets"],
        "riders_by_department": {d: G["by_dept"][d] for d in A.DEPTS},
        "ask_departures": [ampm(d) for d in A.ASK_DEPS],
        "conventions": {"network_releases": [P.RELEASE_2025, P.RELEASE_2026], "standard_clauses": ["2.1", "2.2", "2.3",
                                                                                                 "2.4", "2.5"],
                        "departures_from": ampm(core.DEPARTURES[0]), "departures_to": ampm(core.DEPARTURES[-1]),
                        "epochs": ["5:30", "9:25"]},
        "late_staff_arrivals": [{"departure": ampm(d), "late_weekdays": late_days[d],
                                 "late_staff_arrivals": "{:,}".format(gold["late_arrivals"][d]),
                                 "late_arrival_cost": "{:,}".format(int(round(A.HOLDOVER * gold["late_arrivals"][d]))),
                                 "cost_2027": "{:,}".format(int(round(gold["cost"][d])))} for d in A.ASK_DEPS],
        "cost_inputs": {"day_rate": "%.2f" % A.DAY_RATE, "mile_rate": "%.2f" % A.MILE_RATE,
                        "holdover": "%.2f" % A.HOLDOVER, "service_days_2027": core.working_days_2027(),
                        "deadhead_miles": "%.2f" % G["miles"]["golden"],
                        "deadhead_legs_miles": [A.LEG_MILES_IN, A.LEG_MILES_OUT],
                        "vehicle_charges": "{:,}".format(int(round(G["vehicle_charges"]))),
                        "vehicle_charges_cents": "%.2f" % G["vehicle_charges"]},
        "departures": [{"departure": ampm(d), "on_time": F["golden_counts"][d], "late": late_days[d],
                        "share_pct": "%.1f" % (100.0 * F["golden_counts"][d] / F["N"])} for d in deps],
        "contrasts": {
            "default_path": {"call": ampm(F["R0"]), "on_time": {hm(d): F["R0_counts"][d] for d in deps}},
            "midpoint": {"call": ampm(F["R2"]), "on_time": {hm(d): F["R2_counts"][d] for d in deps}},
            "p90_profile_at_departure": ampm(F["R1"]),
            "p90_profile_walked": ampm(F["R3"]),
            "p50_profile_walked": ampm(F["P50"]),
            "walked_at_default_call": F["walked_715"],
            "grid": {k: (ampm(v) if isinstance(v, (int,)) and k not in ("weekdays_all",) and v >= 360 else v)
                     for k, v in F["grid"].items() if not isinstance(v, dict)},
            "holidays_kept_counts": {hm(d): v for d, v in F["grid"]["holidays_kept_counts"].items()},
            "strict_counts": {hm(d): v for d, v in F["grid"]["strict"].items()},
            "all_vehicle_counts": {hm(d): v for d, v in F["walk_all_counts"].items()},
            "nearest_arrival_seconds": {hm(d): "%.1f" % v for d, v in F["nearest_seconds"].items()},
            "mean_trip_minutes_walked": {hm(d): "%.1f" % v for d, v in F["mean_trip_walk"].items()},
            "mean_trip_minutes_default": {hm(d): "%.1f" % v for d, v in F["mean_trip_R0"].items()},
            "herded_days": len(W["herded"]),
            "herd_factor_range": ["%.2f" % min(W["herded"].values()), "%.2f" % max(W["herded"].values())],
            "incident_days": int(W["incident"][W["work"]].sum()),
            "moved_at_default_call": F["moved_715"],
            "queue_at_default_call": {k: "%.1f" % v for k, v in F["merge_715"].items()},
            "school_term_days": int((W["kind"][W["work"]] == "school").sum()),
            "break_days": int((W["kind"][W["work"]] == "break").sum()),
        },
        "ask_stops": {
            "riders": {k: v for k, v in G["riders"].items()},
            "vehicles": {k: A.vehicles(v) for k, v in G["riders"].items()},
            "seats_empty": {k: A.vehicles(v) * A.SEATS - v for k, v in G["riders"].items()},
            "deadhead_miles": {k: "%.2f" % v for k, v in G["miles"].items()},
            "late_days_default_path": [F["late"]["R0"][d] for d in A.ASK_DEPS],
            "late_days_midpoint": [F["late"]["R2"][d] for d in A.ASK_DEPS],
            "late_days_walked": [F["late"]["walk"][d] for d in A.ASK_DEPS],
            "cracker_cost": ["{:,}".format(int(round(G["stops"][("R1", "walk", "current")]["cost"][d])))
                             for d in A.ASK_DEPS],
            "nearest_composition": {"units": AN.nearest_composition(G)[0],
                                    "subset": "+".join(AN.nearest_composition(G)[1]),
                                    "departure": ampm(AN.nearest_composition(G)[2]),
                                    "figure": AN.nearest_composition(G)[3]},
            "deadhead_cost_gap": {k: "{:,.2f}".format(G["gold"]["cost"][A.ASK_DEPS[0]] -
                                                      G["stops"][("golden", "walk", k)]["cost"][A.ASK_DEPS[0]])
                                  for k in ("current", "drop")},
            "deadhead_cost_gap_pct": {k: ["%.2f" % (100 * (G["gold"]["cost"][d] - G["stops"][("golden", "walk", k)]
                                                         ["cost"][d]) / G["gold"]["cost"][d]) for d in A.ASK_DEPS]
                                      for k in ("current", "drop")},
            "compositions": len(G["stops"]) - 1,
            "cracker_late_arrivals": ["{:,}".format(G["stops"][("R1", "walk", "current")]["late_arrivals"][d])
                                      for d in A.ASK_DEPS],
        },
    }
    return out


def docs_text(target):
    """Every filed sentence: the documents, the readme and every string cell of the pack's workbooks."""
    from openpyxl import load_workbook
    out = {"standard": "\n".join(t for _, t in D.STANDARD), "budget": "\n".join(t for _, t in D.BUDGET),
           "notice": "\n".join(t for _, t in D.NOTICE), "readme": open(os.path.join(target, P.FILES["readme"])).read()}
    for k in ("route", "deadhead", "charges", "holidays", "reassign"):
        wb = load_workbook(os.path.join(target, P.FILES[k]), read_only=True)
        out[k] = "\n".join(str(v) for ws in wb.worksheets for r in ws.iter_rows(values_only=True) for v in r
                           if isinstance(v, str))
    return out


def main():
    W, S, F, G = write_all(TARGET, GOLDEN)
    metadata()
    n, cells = C.run(W, S, F, G, TARGET, GOLDEN, TASK, docs_text(TARGET))
    print("checks: %d assertions green (%d composition cells)" % (n, cells))
    fj = figures_json(W, S, F, G)
    fj["ask_stops"]["composition_cells"] = cells
    with open(os.path.join(HERE, "figures.json"), "w") as fh:
        json.dump(fj, fh, indent=1)
        fh.write("\n")
    v = V.run(TARGET)
    gold = G["gold"]
    assert v["N"] == F["N"] and v["need"] == F["need"] and v["call"] == F["call"], (v["call"], F["call"])
    assert all(v["on_time"][d] == F["golden_counts"][d] for d in core.DEPARTURES), v["on_time"]
    assert v["riders"] == gold["riders"] and v["vehicles"] == gold["vehicles"] and v["seats_empty"] == gold["seats_empty"]
    assert v["both_sheets"] == G["both_sheets"] and v["by_dept"] == {k: x for k, x in G["by_dept"].items() if x}
    assert abs(v["miles"] - G["miles"]["golden"]) < 1e-9 and v["release"] == P.RELEASE_2025
    assert v["service_days"] == core.working_days_2027()
    for d in A.ASK_DEPS:
        assert v["late_arrivals"][d] == gold["late_arrivals"][d]
        assert abs(v["cost"][d] - gold["cost"][d]) < 0.005, (d, v["cost"][d], gold["cost"][d])
    print("verifier: green (call %s, %d riders, nearest arrival to 8:00 %.1f s)" % (ampm(v["call"]), v["riders"],
                                                                                  v["nearest_seconds"]))
    if "--once" not in sys.argv:
        tmp = tempfile.mkdtemp()
        t2, g2 = os.path.join(tmp, "target"), os.path.join(tmp, "golden")
        write_all(t2, g2)
        for a, b in ((TARGET, t2), (GOLDEN, g2)):
            fa, fb = sorted(os.listdir(a)), sorted(os.listdir(b))
            assert fa == fb
            for x in fa:
                assert filecmp.cmp(os.path.join(a, x), os.path.join(b, x), shallow=False), x
        shutil.rmtree(tmp)
        print("second build: byte-identical (%d pack files, %d goldens)" % (len(os.listdir(TARGET)),
                                                                          len(os.listdir(GOLDEN))))
    zp = os.path.join(TASK, "target.zip")
    with zipfile.ZipFile(zp, "w", zipfile.ZIP_DEFLATED) as z:
        for x in sorted(os.listdir(TARGET)):
            zi = zipfile.ZipInfo(x, date_time=(2026, 10, 16, 9, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            with open(os.path.join(TARGET, x), "rb") as fh:
                z.writestr(zi, fh.read())
    print("target.zip cut (%d bytes)" % os.path.getsize(zp))


if __name__ == "__main__":
    main()
