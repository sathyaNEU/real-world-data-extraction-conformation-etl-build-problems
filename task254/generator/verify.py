"""Independent verifier: reads only task254/target/ (and the golden workbook) with its own parsers, its own
course builder and its own reverse-time estimator, and recomputes every graded figure and the ILD ladder."""
import csv
import datetime as dt
import math
import os
import sys
import xml.etree.ElementTree as ET
from collections import defaultdict

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
LOCK = dt.date(2026, 9, 30)
IVL = dt.date(2025, 10, 1)
CONCERN = {"Interstitial lung disease": "ild", "Herpes zoster": "zoster", "Serious infections": "infection",
           "Venous thromboembolism": "vte", "Hepatotoxicity": "hepatic", "Neutropenia": "neutropenia"}
RISKS = ["ild", "zoster", "infection", "vte", "hepatic", "neutropenia"]
NEUT_PER_109 = {"10^9/L": 1.0, "G/l": 1.0, "/nl": 1.0, "10^3/µL": 1.0, "/mm3": 1000.0, "cell/µL": 1000.0, "/µl": 1000.0}


def d(s):
    return dt.date(int(s[:4]), int(s[5:7]), int(s[8:10])) if len(s) == 10 else None


def rows_csv(path):
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def sheet(path, name=None, header_row=1):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[name] if name else wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    for _ in range(header_row - 1):
        next(it)
    hdr = [str(h) for h in next(it)]
    out = [dict(zip(hdr, r)) for r in it if any(v is not None for v in r)]
    wb.close()
    return out


def onset_days(target):
    f = lambda n: os.path.join(target, n)
    lst = rows_csv(f("lyrevo_line_listing_PSUR6_DLP20260930.csv"))
    last = {}
    for r in lst:
        if d(r["version_receipt_date"]) <= LOCK:
            last[r["case_no"]] = max(last.get(r["case_no"], 0), int(r["version_no"]))
    dups = {str(r["duplicate_case_no"]) for r in sheet(f("duplicate_register_PSUR6.xlsx"))}
    terms = {r["preferred_term"]: CONCERN[r["safety_concern"]] for r in sheet(f("rmp_term_list_version_4.xlsx"), "Term list v4")
             if r["safety_concern"] in CONCERN}
    reac = defaultdict(list)
    sex = {}
    for r in lst:
        c = r["case_no"]
        if last.get(c) != int(r["version_no"]) or r["report_type"] == "ST" or c in dups:
            continue
        sex[c] = r["sex"]
        k = terms.get(r["reaction_pt"])
        if k:
            reac[(c, k)].append(r["onset_date"])
    # laboratory onsets
    reg = {}
    for r in sheet(f("psp_laboratory_register.xlsx"), "register"):
        reg[(r["lab_code"], r["analyte"])] = r
    first_q = {}
    for r in rows_csv(f("psp_lab_feed_2023-2026.csv")):
        c, an, val = r["case_no"], r["analyte"], float(r["value"])
        if an == "ALT":
            g = reg[(r["lab_code"], "ALT")]
            uln = float(g["uln_male"] if sex.get(c) == "M" else g["uln_female"])
            hit, key = val > 3 * uln, (c, "hepatic")
        elif an == "NEUT":
            hit, key = val / NEUT_PER_109[r["unit"]] < 1.0, (c, "neutropenia")
        else:
            continue
        if hit:
            sd = d(r["sample_date"])
            if key not in first_q or sd < first_q[key]:
                first_q[key] = sd
    lines = defaultdict(list)
    for r in rows_csv(f("lyrevo_dosage_lines_PSUR6.csv")):
        lines[r["case_no"]].append(r)
    out = defaultdict(list)
    for (c, k), ons in reac.items():
        full = [d(o) for o in ons if len(o) == 10]
        if full:
            on = min(full)
        elif all(o == "" for o in ons) and k in ("hepatic", "neutropenia") and (c, k) in first_q:
            on = first_q[(c, k)]
        else:
            continue
        ln = sorted(lines[c], key=lambda r: int(r["line_no"]))
        if any(len(r["start_date"]) != 10 for r in ln):
            continue
        start = None
        prev_stop = None
        cstart = None
        for r in ln:
            s0 = d(r["start_date"])
            if cstart is None or (prev_stop is not None and (s0 - prev_stop).days - 1 > 28):
                cstart = s0
            if s0 <= on:
                start = cstart
            prev_stop = d(r["stop_date"]) if r["stop_date"] else None
        if start is None:
            continue
        out[k].append(((on - start).days + 1, (LOCK - start).days + 1, start))
    return out


def reverse_time_cdf(pairs):
    """F(t) for right-truncated onsets, built as a product over event days from the largest down."""
    ev = defaultdict(int)
    for x, T, _ in pairs:
        ev[x] += 1
    days = sorted(ev, reverse=True)
    F = {}
    surv = 1.0
    for s in days:
        F[s] = surv
        risk = sum(1 for x, T, _ in pairs if x <= s <= T)
        surv *= 1.0 - ev[s] / risk
    return sorted(F.items())


def month_cut(Fs, p):
    for m in range(1, 60):
        val = 0.0
        for s, f in Fs:
            if s <= 30 * m:
                val = f
        if val >= p - 1e-12:
            return m


def naive_cut(xs, p):
    xs = sorted(xs)
    return math.ceil(xs[math.ceil(p * len(xs)) - 1] / 30.0)


def exposure(target):
    f = lambda n: os.path.join(target, n)
    root = ET.parse(f("xevmpd_presentations_RVB-417.xml")).getroot()
    ns = root.tag.split("}")[0] + "}"
    pres = {}
    for ap in root.iter(ns + "authorisedProduct"):
        pres[ap.find(ns + "authorisationNumber").text] = (int(ap.find(ns + "strength").get("numerator")),
                                                         int(ap.find(ns + "package/" + ns + "packSize").text))
    code = {}
    for r in rows_csv(f("product_master_national_codes.csv")):
        code[(r["market"], r["national_code"].lstrip("0"))] = r["eu_presentation"]
    tot = defaultdict(float)
    for r in rows_csv(f("sellin_wholesaler_monthly_2023-2026.csv")):
        if r["posting_period"] > "2026-09":
            continue
        st, ps = pres[code[(r["country"], r["national_code"].lstrip("0"))]]
        tabs = int(r["quantity"]) * ps
        tot["launch"] += tabs
        tot[st] += tabs
        if r["posting_period"] >= "2025-10":
            tot["interval"] += tabs
    wb = openpyxl.load_workbook(f("tender_calloffs_FR_IT_ES_2025-2026.xlsx"), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    hdr = None
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == "call_off_no":
            hdr = list(row)
            continue
        if hdr is None or row[0] is None:
            continue
        r = dict(zip(hdr, row))
        dd = r["delivery_date"].date() if hasattr(r["delivery_date"], "date") else r["delivery_date"]
        if dd > LOCK:
            continue
        st, ps = pres[code[(r["country"], str(r["gtin"]).lstrip("0"))]]
        tabs = int(r["qty"]) * 10 * ps
        tot["launch"] += tabs
        tot[st] += tabs
        if dd >= IVL:
            tot["interval"] += tabs
    wb.close()
    return {k: v / 365.25 for k, v in tot.items()}


def counts(target):
    f = lambda n: os.path.join(target, n)
    lst = rows_csv(f("lyrevo_line_listing_PSUR6_DLP20260930.csv"))
    dups = {str(r["duplicate_case_no"]) for r in sheet(f("duplicate_register_PSUR6.xlsx"))}
    terms = {r["preferred_term"]: CONCERN[r["safety_concern"]] for r in sheet(f("rmp_term_list_version_4.xlsx"), "Term list v4")
             if r["safety_concern"] in CONCERN}
    last = {}
    for r in lst:
        if d(r["version_receipt_date"]) <= LOCK:
            last[r["case_no"]] = max(last.get(r["case_no"], 0), int(r["version_no"]))
    seen = {k: set() for k in RISKS}
    ivl = {k: set() for k in RISKS}
    for r in lst:
        c = r["case_no"]
        if last.get(c) != int(r["version_no"]) or r["report_type"] == "ST" or c in dups:
            continue
        k = terms.get(r["reaction_pt"])
        if k:
            seen[k].add(c)
            if d(r["initial_receipt_date"]) >= IVL:
                ivl[k].add(c)
    return {k: {"launch": len(seen[k]), "interval": len(ivl[k])} for k in RISKS}


def main(target=None, golden_dir=None):
    target = target or os.path.join(TASK, "target")
    golden_dir = golden_dir or os.path.join(TASK, "golden")
    on = onset_days(target)
    res = {"windows": {}, "naive": {}}
    for k in RISKS:
        Fs = reverse_time_cdf(on[k])
        res["windows"][k] = [month_cut(Fs, 0.8), month_cut(Fs, 0.95)]
        xs = [x for x, _, _ in on[k]]
        res["naive"][k] = [naive_cut(xs, 0.8), naive_cut(xs, 0.95)]
    ild = on["ild"]
    res["ild_rungs"] = {
        "drop2026": [naive_cut([x for x, _, s in ild if s < dt.date(2026, 1, 1)], p) for p in (0.8, 0.95)],
        "year": [naive_cut([x for x, _, s in ild if s <= dt.date(2025, 9, 30)], p) for p in (0.8, 0.95)],
    }
    ex = exposure(target)
    cn = counts(target)
    res["py"] = {k: int(round(ex[k] / 10.0)) * 10 for k in ("launch", "interval", 15, 30)}
    res["rates"] = {k: {w: round(1000.0 * cn[k][w] / ex[w] + 1e-9, 1) for w in ("launch", "interval")} for k in RISKS}
    # the golden workbook says the same
    gw = os.path.join(golden_dir, "rmp_monitoring_windows_2027.xlsx")
    if os.path.exists(gw):
        wb = openpyxl.load_workbook(gw, data_only=True)
        found = 0
        names = {"Interstitial lung disease": "ild", "Herpes zoster": "zoster", "Serious infections": "infection",
                 "Venous thromboembolism": "vte", "Hepatotoxicity": "hepatic", "Neutropenia": "neutropenia"}
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                if row and row[0] in names:
                    k = names[row[0]]
                    nums = [v for v in row[1:] if isinstance(v, (int, float))]
                    if ws.title.lower().startswith("monitoring") and len(nums) >= 2:
                        assert [int(nums[0]), int(nums[1])] == res["windows"][k], (k, nums, res["windows"][k])
                        found += 1
                    if ws.title.lower().startswith("frequency") and len(nums) >= 2:
                        assert abs(nums[0] - res["rates"][k]["launch"]) < 1e-9, (k, nums, res["rates"][k])
                        assert abs(nums[1] - res["rates"][k]["interval"]) < 1e-9, (k, nums, res["rates"][k])
                        found += 1
        assert found == 12, found
        res["golden_workbook_rows_checked"] = found
    return res


if __name__ == "__main__":
    out = main(*(sys.argv[1:3]))
    for k, v in out.items():
        print(k, v)
