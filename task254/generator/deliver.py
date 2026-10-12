"""Write the golden deliverables into task254/golden/ from the shipped bundle alone."""
import os
import sys

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import analysis as A
import estimators as E
import golden as G
import postprocess as PP

TASK = os.path.dirname(HERE)
LABEL = {"ild": "Interstitial lung disease", "zoster": "Herpes zoster", "infection": "Serious infections",
         "vte": "Venous thromboembolism", "hepatic": "Hepatotoxicity", "neutropenia": "Neutropenia"}
ORDER = ["ild", "zoster", "infection", "vte", "hepatic", "neutropenia"]
INK, INK2, MUTED, GRID, SURFACE = "#1d1d1b", "#52514e", "#8a8984", "#e2e1dc", "#ffffff"
BLUE, ILD_C = "#2a78d6", "#c2410c"


def monthly(pack, r, months=24):
    t = A.tto(pack, r)
    s, F = E.reverse_km(t["x"].to_numpy(float), t["T"].to_numpy(float))
    return [E.cdf_at(s, F, 30 * m) for m in range(0, months + 1)]


def workbook(path, g):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Monitoring windows"
    bold = Font(bold=True)
    hfill = PatternFill("solid", fgColor="D9E1EA")
    thin = Side(style="thin", color="A6A6A6")
    ild95 = g["windows"]["ild"][1]
    ws["A1"] = "ILD monitoring runs to month %d: monitoring windows for the six identified risks" % ild95
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Lyrevo (RVB-417), RMP update filed with PSUR 6 (data lock point 30 Sep 2026). Prepared by Drug Safety, 7 Oct 2026."
    ws["A4"] = ("For RMP Part V: pulmonary monitoring for patients starting treatment from January 2027 ends after month %d "
                "of treatment." % ild95)
    ws["A4"].font = Font(bold=True, color="9A3412")
    hdr = ["Identified risk", "80% of reactions begun by month", "95% of reactions begun by month"]
    for j, h in enumerate(hdr):
        c = ws.cell(row=6, column=j + 1, value=h)
        c.font = bold
        c.fill = hfill
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(bottom=thin)
    for i, r in enumerate(ORDER):
        w = g["windows"][r]
        ws.cell(row=7 + i, column=1, value=LABEL[r])
        ws.cell(row=7 + i, column=2, value=int(w[0]))
        ws.cell(row=7 + i, column=3, value=int(w[1]))
        if r == "ild":
            for j in range(1, 4):
                ws.cell(row=7 + i, column=j).font = bold
    ws["A14"] = ("Source: Argus line listing at the DLP, dosage lines and PSP laboratory feed (RMP annex 7). Each case is "
                 "bounded by its own time on treatment at the lock, the days from the first dose of its course to 30 Sep "
                 "2026, in a reverse-time product-limit estimate; no parametric model.")
    ws["A14"].font = Font(italic=True, color="595959", size=9)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 20
    ws.row_dimensions[6].height = 32
    for row in ws.iter_rows(min_row=7, max_row=12, min_col=2, max_col=3):
        for c in row:
            c.alignment = Alignment(horizontal="right")
            c.number_format = "0"
    ws.freeze_panes = "A7"
    ws.print_area = "A1:C14"
    ws.print_title_rows = "6:6"
    ws.page_setup.orientation = "landscape"

    fs = wb.create_sheet("Frequency")
    fs["A1"] = "Reporting rates per 1,000 patient-years and the exposure under them"
    fs["A1"].font = Font(bold=True, size=12)
    fs["A2"] = ("Cases as they stood at the DLP (interval: first received 1 Oct 2025 to 30 Sep 2026). Exposure from units "
                "supplied through every channel (PSUR section 5), one tablet one patient-day, 365.25 days a year.")
    hdr = ["Identified risk", "Rate from launch (3 Apr 2023 to 30 Sep 2026)", "Rate, 1 Oct 2025 to 30 Sep 2026"]
    for j, h in enumerate(hdr):
        c = fs.cell(row=4, column=j + 1, value=h)
        c.font = bold
        c.fill = hfill
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(bottom=thin)
    for i, r in enumerate(ORDER):
        fs.cell(row=5 + i, column=1, value=LABEL[r])
        for j, w in enumerate(("launch", "interval")):
            c = fs.cell(row=5 + i, column=2 + j, value=round(g["rates"][r][w] + 1e-9, 1))
            c.number_format = "0.0"
            c.alignment = Alignment(horizontal="right")
    r10 = lambda v: int(round(v / 10.0)) * 10
    fs["A12"] = "Patient-years (nearest 10)"
    fs["A12"].font = bold
    fs["A12"].fill = hfill
    rows = [("From launch to 30 Sep 2026", r10(g["py"]["launch"])),
            ("   of which 15 mg tablets", r10(g["py_strength"][15])),
            ("   of which 30 mg tablets", r10(g["py_strength"][30])),
            ("1 Oct 2025 to 30 Sep 2026", r10(g["py"]["interval"]))]
    for i, (lab, v) in enumerate(rows):
        fs.cell(row=13 + i, column=1, value=lab)
        c = fs.cell(row=13 + i, column=2, value=v)
        c.number_format = "#,##0"
        c.alignment = Alignment(horizontal="right")
    fs.column_dimensions["A"].width = 30
    fs.column_dimensions["B"].width = 24
    fs.column_dimensions["C"].width = 22
    fs.row_dimensions[4].height = 32
    fs.freeze_panes = "A5"
    fs.print_area = "A1:C16"
    fs.page_setup.orientation = "landscape"
    ns = wb.create_sheet("Notes")
    notes = [
        ("Definitions", None),
        ("Case", "Post-marketing case (spontaneous, solicited, literature) as it stood at the DLP, the latest version received "
                 "on or before 30 Sep 2026; duplicates as one case under the master of the Case Processing register; "
                 "registry and other study cases not used (annex 7, sections 1 and 2)."),
        ("Identified risk", "Reactions mapped by RMP term list v4; one reaction per case and risk, the earliest onset."),
        ("Time to onset", "Days from the first dose of the course in which the reaction began (day 1); a course ends at a "
                          "gap of more than 28 days without a dose; complete dates only (annex 7, section 3)."),
        ("Laboratory-only reactions", "Hepatotoxicity and neutropenia without a reporter onset date: date of the first "
                                      "PSP sample meeting the case definition (annex 7, section 4)."),
        ("Window", "First 30-day month of treatment by the end of which at least 80 or 95 per cent of the risk's "
                   "reactions have begun (annex 7, section 5)."),
        ("Reporting rate", "Cases per 1,000 patient-years; for the interval, cases first received 1 Oct 2025 to "
                           "30 Sep 2026 (annex 7, section 6)."),
        ("Patient-years", "Units supplied up to the DLP through wholesalers and hospital call-offs, returns netted, free "
                          "packs counted, one tablet one patient-day, 365.25 days (PSUR section 5); nearest 10."),
    ]
    for i, (k, t) in enumerate(notes):
        a = ns.cell(row=1 + i, column=1, value=k)
        a.font = bold if i == 0 or t else Font()
        if t:
            b = ns.cell(row=1 + i, column=2, value=t)
            b.alignment = Alignment(wrap_text=True, vertical="top")
            a.alignment = Alignment(vertical="top")
    ns.column_dimensions["A"].width = 24
    ns.column_dimensions["B"].width = 100
    wb.properties.creator = "Lasse Christoffersen"
    wb.properties.lastModifiedBy = "Lasse Christoffersen"
    import datetime as dt
    wb.properties.created = dt.datetime(2026, 10, 6, 13, 0)
    wb.properties.modified = dt.datetime(2026, 10, 7, 10, 24)
    wb.save(path)
    PP.normalize_zip(path, modified="2026-10-07T10:24:00Z")


def chart(path, pack, g):
    plt.rcParams.update({"font.family": "Liberation Sans", "font.size": 8.5, "axes.edgecolor": MUTED,
                         "xtick.color": INK2, "ytick.color": INK2, "axes.labelcolor": INK2,
                         "svg.fonttype": "none", "svg.hashsalt": "task254"})
    fig, axes = plt.subplots(2, 3, figsize=(10.2, 6.2), dpi=100, sharex=True, sharey=True)
    fig.patch.set_facecolor(SURFACE)
    months = np.arange(0, 25)
    for ax, r in zip(axes.flat, ORDER):
        ax.set_facecolor(SURFACE)
        cum = np.array(monthly(pack, r)) * 100
        col = ILD_C if r == "ild" else BLUE
        ax.step(months, cum, where="pre", color=col, lw=2, solid_joinstyle="miter")
        ax.fill_between(months, cum, step="pre", color=col, alpha=0.10, lw=0)
        w80, w95 = g["windows"][r]
        ax.axhline(95, color=MUTED, lw=0.8, ls=(0, (2, 2)), zorder=1)
        ax.axvline(w95, color=INK, lw=1.1, ls=(0, (4, 2)), zorder=3)
        ax.plot([w95], [cum[w95]], marker="o", ms=5.5, color=col, markeredgecolor=SURFACE, mew=1.5, zorder=4)
        right = w95 <= 12
        off, ha = ((6, "left") if right else (-5, "right"))
        ax.annotate("95%%: month %d" % w95, (w95, 50), xytext=(off, 0), textcoords="offset points", ha=ha,
                    fontsize=8.5, color=INK, fontweight="bold" if r == "ild" else "normal")
        ax.annotate("80%%: month %d" % w80, (w95, 40), xytext=(off, 0), textcoords="offset points", ha=ha,
                    fontsize=7.5, color=INK2)
        ax.set_title(LABEL[r], loc="left", fontsize=9.5, color=INK, fontweight="bold" if r == "ild" else "normal")
        ax.set_xlim(0, 24)
        ax.set_ylim(0, 101)
        ax.set_xticks([0, 6, 12, 18, 24])
        ax.set_yticks([0, 25, 50, 75, 95])
        ax.grid(axis="y", color=GRID, lw=0.6)
        ax.set_axisbelow(True)
        for s in ("top", "right"):
            ax.spines[s].set_visible(False)
    for ax in axes[1]:
        ax.set_xlabel("month of treatment")
    for ax in axes[:, 0]:
        ax.set_ylabel("reactions begun, cumulative %")
    ild95 = g["windows"]["ild"][1]
    fig.suptitle("Pulmonary monitoring to month %d: by then 95%% of interstitial lung disease reactions have begun" % ild95,
                 x=0.012, ha="left", fontsize=12, fontweight="bold", color=INK)
    fig.text(0.012, 0.905, "Lyrevo, identified risks. Share of each risk's reactions begun by the end of each 30-day month "
                           "of treatment, post-marketing cases at the 30 Sep 2026 data lock,\neach case bounded by its own "
                           "time on treatment at the lock. Dashed line: the 95% window.", fontsize=8, color=INK2, ha="left",
             linespacing=1.4)
    fig.tight_layout(rect=(0, 0, 1, 0.915))
    fig.savefig(path, format="svg", metadata={"Date": None, "Creator": None}, facecolor=SURFACE)
    plt.close(fig)
    # laid out in Liberation Sans (Arial's metrics); the file names the font stack a reader's machine has
    svg = open(path, encoding="utf-8").read()
    svg = svg.replace("'Liberation Sans'", "Arial, 'Liberation Sans', Helvetica, sans-serif")
    svg = svg.replace('id="matplotlib.axis_', 'id="axis_')
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(svg)


def main(target=None, out=None):
    target = target or os.path.join(TASK, "target")
    out = out or os.path.join(TASK, "golden")
    os.makedirs(out, exist_ok=True)
    pack, sup, g = G.solve(target)
    workbook(os.path.join(out, "rmp_monitoring_windows_2027.xlsx"), g)
    chart(os.path.join(out, "onset_windows_by_risk.svg"), pack, g)
    return g


if __name__ == "__main__":
    main()
