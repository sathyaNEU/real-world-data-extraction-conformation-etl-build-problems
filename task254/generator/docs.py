"""The pack's documents: the shuttle service standard and the finance budget note (DOCX), the garage closure
notice and the express route timetable (PDF), and the vendor's readme (TXT)."""
import datetime as dt
import io
import re
import zipfile

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import asks as A
import postprocess as PP

ORG = "Beargrass Medical Center"

STANDARD = [
    ("t", "Staff Shuttle Service Standard"),
    ("s", "Transportation Planning, version 2, June 2026. Owner: Ebony Robinson, Transportation Planning Lead."),
    ("h", "1  Purpose"),
    ("p", "1.1  This standard sets out how Transportation Planning sets and reviews the timetables of the "
          "staff shuttles Beargrass Medical Center contracts between its campuses and its remote parking."),
    ("p", "1.2  It covers the timing of departures. Fares, accessibility, vehicle specification and the "
          "contract with the operator are covered by the service order and the operator's schedule of charges."),
    ("h", "2  Setting a departure"),
    ("p", "2.1  A shuttle departure serving a shift start is set so that its trips would have reached the campus "
          "stop by the shift start on at least 90 per cent of the working weekdays in the latest twelve months of "
          "probe travel-time data."),
    ("p", "2.2  A working weekday is a Monday to Friday that is not a hospital holiday in the HR holiday calendar."),
    ("p", "2.3  Every working weekday counts, whatever happened on the road that day."),
    ("p", "2.4  A trip runs from the departure at the boarding bay to arrival at the campus stop, along the route in "
          "the route file, with the route file's fixed legs."),
    ("p", "2.5  Trip times use the passenger-vehicle travel times in the vendor's export."),
    ("p", "2.6  Departures are set on the five-minute grid."),
    ("h", "3  Review"),
    ("p", "3.1  Timetables are reviewed each year when the latest twelve months of probe data are delivered, and "
          "when a garage, lot or route changes."),
    ("p", "3.2  The operator is given the departure times with the service order; running times inside the "
          "timetable are the operator's to set."),
    ("h", "4  Records"),
    ("p", "4.1  Transportation Planning keeps the route files, the probe data and the review workings for three "
          "years."),
]

BUDGET = [
    ("t", "Staff shuttle service: 2027 budget basis"),
    ("s", "Finance, Operating Budget. 12 October 2026. Prepared with Transportation Planning for the 2027 service "
          "order with Ridgeline Transit Services."),
    ("h", "Purpose"),
    ("p", "The Chestnut Street garage closes on 31 December 2026 and its permit holders park at Cedar Grove from 4 "
          "January 2027. This note sets the basis on which Finance budgets the contracted morning shuttle from Cedar "
          "Grove to the Brook Street stop for 2027, so that every departure Transportation Planning considers for the "
          "service order is costed the same way. The shuttle is charged to cost centre 7140 Parking and "
          "Transportation, account 64210 Contracted transport services."),
    ("h", "Service days and vehicles"),
    ("p", "The 2027 service runs Monday to Friday, less the hospital's observed holidays in the HR 2027 calendar."),
    ("p", "The morning departure runs as one convoy with a seat for every rider, at the seats per vehicle in the "
          "operator's 2027 schedule of charges. Standees are not budgeted."),
    ("h", "Operator charges"),
    ("p", "Each vehicle is charged the operator's day charge for the morning run and, at the deadhead rate, the "
          "miles from the Brook Street stop back to the operator's yard along the operator's deadhead route. Rates "
          "are those in the 2027 schedule of charges; the early-run differential and cancellation charges are not "
          "budgeted."),
    ("h", "Late arrivals"),
    ("p", "Late arrivals are budgeted from the latest twelve months of probe data, taken as 2027's: each working "
          "weekday on which a departure's trips would not have reached the stop by 8:00 a.m. counts as a late "
          "arrival for every rider. Each late staff arrival is costed at $%.2f, the hospital's average hold-over and "
          "lost-time cost for 2026." % A.HOLDOVER),
    ("h", "Cost of holding a departure"),
    ("p", "The cost of holding a departure in 2027 is the vehicles' charges for the year plus the cost of its late "
          "arrivals. Amounts are carried in dollars and cents and rounded to whole dollars in the budget."),
    ("p", "Contact: Nicholas King, Operating Budget, extension 4-3108."),
]

NOTICE = [
    ("t", "Chestnut Street Garage closure: staff parking from 4 January 2027"),
    ("s", "Parking Services, 15 September 2026"),
    ("p", "The Chestnut Street garage closes after the evening shift on Thursday 31 December 2026 so that it can "
          "be demolished for the new outpatient pavilion. Demolition and construction will take about two years."),
    ("p", "From Monday 4 January 2027, staff who hold a Chestnut Street permit are reassigned to the Cedar Grove "
          "park-and-ride lot on KY-480, about twenty miles south of campus. A contracted staff shuttle runs from "
          "Cedar Grove to the Brook Street stop."),
    ("p", "The morning shuttle is timed for staff whose scheduled start is 8:00 a.m. Each reassigned member of "
          "staff with an 8:00 a.m. start has a seat on it. Staff with other start times will be told of their "
          "runs in December."),
    ("p", "Staff who would rather not use Cedar Grove can opt out through Parking Services and keep a transit "
          "pass or a carpool space instead. Staff without a Chestnut Street permit can ask for a Cedar Grove space "
          "through the HR Service Center; a space granted that way carries the same seat on the shuttle."),
    ("p", "Questions: Parking Services, extension 4-2275 (Andrew Brown, Parking Services Manager)."),
]


def _styles():
    ss = getSampleStyleSheet()
    return dict(
        body=ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=10.2, leading=14, spaceAfter=8),
        title=ParagraphStyle("t", parent=ss["Normal"], fontName="Helvetica-Bold", fontSize=15, leading=19,
                             spaceAfter=6),
        small=ParagraphStyle("s", parent=ss["Normal"], fontName="Helvetica", fontSize=8.6, leading=11,
                             textColor=colors.HexColor("#444444"), spaceAfter=12),
        cell=ParagraphStyle("c", parent=ss["Normal"], fontName="Helvetica", fontSize=8.6, leading=10.5),
    )


def _footer(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(0.8 * inch, 0.55 * inch, text)
        canvas.drawRightString(LETTER[0] - 0.8 * inch, 0.55 * inch, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def notice_pdf(path):
    st = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=LETTER, leftMargin=0.9 * inch, rightMargin=0.9 * inch,
                            topMargin=0.9 * inch, bottomMargin=0.9 * inch, title="Chestnut Street Garage closure",
                            author="Parking Services", creator="Parking Services", producer="Parking Services")
    story = []
    for kind, text in NOTICE:
        story.append(Paragraph(text, st["title"] if kind == "t" else st["small"] if kind == "s" else st["body"]))
    doc.build(story, onFirstPage=_footer("%s  |  Parking Services" % ORG),
              onLaterPages=_footer("%s  |  Parking Services" % ORG))
    open(path, "wb").write(buf.getvalue())
    PP.normalize_pdf(path, b"D:20260915141200-04'00'", "Parking Services")


def transit_pdf(path):
    """A neighbouring file: the regional transit authority's Route 40X express timetable (another corridor)."""
    st = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=LETTER, leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                            topMargin=0.8 * inch, bottomMargin=0.8 * inch, title="Route 40X Southwest Express",
                            author="Regional transit authority", creator="Regional transit authority",
                            producer="Regional transit authority")
    stops = ["Valley Station P&R", "US-31W at Greenwood Rd", "US-31W at Algonquin Pkwy",
             "W Broadway at 18th St", "4th St at Broadway", "Medical Center (Chestnut St)"]
    runs = []
    t0 = 5 * 60 + 40
    for k in range(9):
        dep = t0 + 15 * k
        seg = [0, 9, 18, 31, 37, 42] if k < 3 else [0, 10, 21, 36, 43, 48] if k < 7 else [0, 9, 19, 33, 40, 45]
        runs.append(["%d:%02d" % ((dep + s) // 60, (dep + s) % 60) for s in seg])
    data = [[Paragraph(x, st["cell"]) for x in stops]] + [list(r) for r in runs]
    tb = Table(data, repeatRows=1)
    tb.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF4")),
                            ("FONTSIZE", (0, 1), (-1, -1), 8.6), ("ALIGN", (0, 1), (-1, -1), "CENTER")]))
    story = [Paragraph("Route 40X Southwest Express: weekday mornings, inbound", st["title"]),
             Paragraph("Effective 17 August 2026. Times are scheduled; allow extra time in bad weather.", st["small"]),
             tb, Spacer(1, 10),
             Paragraph("Weekday service only. No service on New Year's Day, Memorial Day, Independence Day, Labor Day, "
                       "Thanksgiving Day or Christmas Day.", st["small"])]
    doc.build(story)
    open(path, "wb").write(buf.getvalue())
    PP.normalize_pdf(path, b"D:20260803093000-04'00'", "Regional transit authority")


def _docx(path, content, author, title, created, modified):
    from docx import Document
    from docx.shared import Pt
    d = Document()
    stl = d.styles["Normal"]
    stl.font.name = "Calibri"
    stl.font.size = Pt(11)
    for kind, text in content:
        if kind == "t":
            d.add_heading(text, level=0)
        elif kind == "s":
            p = d.add_paragraph(text)
            p.runs[0].italic = True
        elif kind == "h":
            d.add_heading(text, level=1)
        else:
            d.add_paragraph(text)
    cp = d.core_properties
    cp.author = author
    cp.last_modified_by = author
    cp.title = title
    cp.comments = ""
    cp.keywords = ""
    cp.subject = ""
    cp.category = ""
    cp.revision = 4
    cp.created = created
    cp.modified = modified
    buf = io.BytesIO()
    d.save(buf)
    normalize_docx(buf.getvalue(), path)


def standard_docx(path):
    _docx(path, STANDARD, "Ebony Robinson", "Staff Shuttle Service Standard", dt.datetime(2024, 3, 4, 10, 15),
          dt.datetime(2026, 6, 22, 16, 5))


def budget_docx(path):
    _docx(path, BUDGET, "Nicholas King", "Staff shuttle service: 2027 budget basis", dt.datetime(2026, 10, 7, 9, 40),
          dt.datetime(2026, 10, 12, 17, 26))


def normalize_docx(data, path):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        parts = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    drop = {"word/stylesWithEffects.xml", "docProps/thumbnail.jpeg"}
    drop |= {n for n, _ in parts if n.startswith("customXml/")}
    out = []
    for name, b in parts:
        if name in drop:
            continue
        if name.endswith(".rels") or name == "[Content_Types].xml":
            b = re.sub(rb'<Relationship [^>]*Target="[^"]*(stylesWithEffects\.xml|thumbnail\.jpeg|customXml/[^"]*)"/>',
                       b"", b)
            b = re.sub(rb'<Override PartName="/(word/stylesWithEffects\.xml|docProps/thumbnail\.jpeg|customXml/[^"]*)"'
                       rb'[^>]*/>', b"", b)
            b = re.sub(rb'<Default Extension="jpeg"[^>]*/>', b"", b)
        if name.endswith((".xml", ".rels")):
            b = b.replace(b"<?xml version='1.0' encoding='UTF-8' standalone='yes'?>\n", PP.DECL)
        if name == "docProps/app.xml":
            b = re.sub(rb"<Application>[^<]*</Application>", b"<Application>Microsoft Office Word</Application>", b)
            b = re.sub(rb"<AppVersion>[^<]*</AppVersion>", b"<AppVersion>16.0000</AppVersion>", b)
            b = re.sub(rb"<Company>[^<]*</Company>", b"", b)
        if name == "docProps/core.xml":
            b = re.sub(rb"<cp:lastPrinted>[^<]*</cp:lastPrinted>", b"", b)
            b = b.replace(b"python-docx", b"")
        out.append((name, b))
    with zipfile.ZipFile(path, "w") as z:
        for name, b in out:
            zi = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, b)


def readme_txt(path, files):
    lines = [
        "Trelvane Mobility Data: TMC travel-time export",
        "Customer: Beargrass Medical Center, Transportation Planning (order 2026-0418)",
        "Delivered 5 October 2026",
        "",
        "CONTENTS",
        "",
        "  One CSV per calendar month, October 2025 to September 2026:",
        "  %s to %s" % (files["probe_2025-10"], files["probe_2026-09"]),
        "  %s, the segments in each network release used in the export" % files["tmc"],
        "",
        "  Coverage: I-65 between Exit 121 and Exit 136 and the I-264 interchange, both directions, KY-480 at",
        "  Exit 121; every day; epochs from 05:30 to 09:25.",
        "",
        "FIELDS",
        "",
        "  tmc_code              segment code",
        "  measurement_tstamp    start of the five-minute epoch, local time (Eastern)",
        "  speed_all             average speed, all vehicles, mph",
        "  travel_time_all       average travel time over the segment, all vehicles, seconds",
        "  speed_passenger       average speed, passenger vehicles, mph",
        "  travel_time_passenger average travel time over the segment, passenger vehicles, seconds",
        "  speed_freight         average speed, freight trucks, mph",
        "  travel_time_freight   average travel time over the segment, freight trucks, seconds",
        "  reference_speed       the segment's reference speed (85th percentile off-peak), mph",
        "  data_density          A: 5 or more probe vehicles in the epoch; B: 2 to 4; C: fewer than 2, the",
        "                        value filled from the segment's history for that epoch",
        "",
        "  Travel times are averages over the class's probe vehicles in the epoch. Speeds are length over",
        "  travel time. An epoch with no probe vehicles and no usable history for the segment is left out",
        "  of the export.",
        "",
        "NETWORK RELEASES",
        "",
        "  %s lists every segment of each network release used in the export, with its" % files["tmc"],
        "  length in miles. The release in force on a date is the latest whose effective date is on or",
        "  before it; release 2025.2 was in force to 31 March 2026 and release 2026.1 from 1 April 2026.",
        "",
        "LICENCE",
        "",
        "  For the customer's internal planning use. Not for redistribution.",
        "",
    ]
    with open(path, "w", newline="") as f:
        f.write("\r\n".join(lines))
