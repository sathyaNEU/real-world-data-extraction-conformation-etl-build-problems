"""Every assertion of the build, computed from the files in target/ through the golden's own code path."""
import datetime as dt
import itertools
import json
import os
import re
import shutil
import tempfile

import numpy as np
import pandas as pd

import analysis as A
import core as C
import estimators as E
import golden as G
import pack as P

LOCK = pd.Timestamp(2026, 9, 30)
METHODS = ["inverted_cdf", "linear", "hazen", "weibull", "median_unbiased", "normal_unbiased",
           "interpolated_inverted_cdf", "nearest", "lower", "higher", "midpoint", "averaged_inverted_cdf",
           "closest_observation"]
RUNG = {"ild": {"weibull": [8, 13], "naive": [7, 16], "drop2026": [8, 18], "year": [8, 19]}}
NAIVE = {"ild": [7, 16], "zoster": [8, 16], "infection": [5, 11], "vte": [5, 11], "hepatic": [3, 8], "neutropenia": [2, 7]}
# Ask A ladders: stop -> (estimator, tto switches)
LADDER = {
    "zoster": [("natural read", "naive", dict(origin="listing")), ("half device", "naive", {}),
               ("right method, wrong origin", "lb", dict(origin="listing")),
               ("over-corrected", "lb", dict(origin="line")), ("latest delivered version", "lb", dict(version="latest"))],
    "infection": [("natural read", "naive", dict(origin="listing")), ("half device", "naive", {}),
                  ("right method, wrong origin", "lb", dict(origin="listing")),
                  ("over-corrected", "lb", dict(origin="line")), ("latest delivered version", "lb", dict(version="latest"))],
    "vte": [("natural read", "naive", {}), ("half device", "naive", dict(origin="line")),
            ("over-corrected", "lb", dict(origin="line")), ("latest delivered version", "lb", dict(version="latest"))],
    "hepatic": [("natural read", "naive", dict(lab="drop")), ("half device", "naive", {}),
                ("right method, blank onsets dropped", "lb", dict(lab="drop")),
                ("over-corrected", "lb", dict(lab="grade1")), ("latest delivered version", "lb", dict(version="latest"))],
    "neutropenia": [("natural read", "naive", dict(lab="drop")), ("half device", "naive", {}),
                    ("right method, units ignored", "lb", dict(lab="units")),
                    ("over-corrected", "lb", dict(lab="grade1")), ("latest delivered version", "lb", dict(version="latest"))],
}
DEVICES = {"zoster": [dict(origin="listing"), dict(origin="line")], "infection": [dict(origin="listing"), dict(origin="line")],
           "vte": [dict(origin="line")], "hepatic": [dict(lab="drop"), dict(lab="grade1")],
           "neutropenia": [dict(lab="drop"), dict(lab="units"), dict(lab="grade1")]}
# the one stop that lands on its answer: a neutropenia reaction dated by its first grade-1 sample moves earlier, but
# no PSP sample is taken before day 85, so the 80 per cent cut (month 3) cannot move
KNOWN_SAME = {("neutropenia", 0, "lb", (("lab", "grade1"),), "lock")}
SUPPLY_DEV = {"B-D1 hospital channel": dict(channels="sellin"), "B-H1 pack size": dict(pack="all28"),
              "B-H2 shipper cases": dict(cases=False), "B-H3 returns gross": dict(returns="gross"),
              "B-H3 free goods dropped": dict(foc="drop"), "B-H4 supply after the lock": dict(cutoff="export"),
              "B-H5 mg per DDD": dict(dose="ddd")}
R10 = lambda v: int(round(v / 10.0)) * 10
R1 = lambda v: round(v + 1e-9, 1)


class Check:
    def __init__(self):
        self.n = 0
        self.rec = {}

    def ok(self, name, cond, info=""):
        if not cond:
            raise AssertionError("%s %s" % (name, info))
        self.n += 1

    def eq(self, name, got, want):
        self.ok(name, got == want, "got %r want %r" % (got, want))

    def record(self, name, value):
        self.rec[name] = value


def _mature_all(ck, tag, t, tm, methods):
    x = t["x"].to_numpy(dtype=float)
    T = t["T"].to_numpy(dtype=float)
    st = t["start"]
    for L in range(int(x.max()), 914):
        sub = x[T >= L]
        for mth in methods:
            got = [E.month_of(E.np_quantile(sub, p, mth)) for p in (0.8, 0.95)]
            ck.eq("%s mature L=%d %s" % (tag, L, mth), got, tm)
    sub = x[(st <= pd.Timestamp(2023, 12, 31)).to_numpy()]
    for mth in methods:
        got = [E.month_of(E.np_quantile(sub, p, mth)) for p in (0.8, 0.95)]
        ck.eq("%s courses started in 2023 %s" % (tag, mth), got, tm)


def run(target, world=None, quick=False):
    ck = Check()
    pack, sup, g = G.solve(target)
    files = sorted(os.listdir(target))
    # ------------------------------------------------------------ input gates
    ck.eq("file count", len(files), len(P.FILES))
    ck.ok("ten or more files", len(files) >= 10)
    fmts = {f.rsplit(".", 1)[1] for f in files}
    ck.ok("three or more formats", len(fmts) >= 3, fmts)
    rows = {f: sum(1 for _ in open(os.path.join(target, f))) - 1 for f in files if f.endswith(".csv")}
    ck.ok("a file of 25,000 rows", max(rows.values()) >= 25000, rows)
    ck.ok("listing is the largest file", max(rows, key=rows.get) == P.FILES["listing"])
    ck.ok("two or more distractors", len(P.DISTRACTORS) >= 2)
    for f in files:
        ck.ok("no trap word in file name %s" % f, not re.search(r"distract|trap|truncat|decoy|answer|golden|known_issue", f, re.I))
    for f in files:
        if f.endswith((".csv", ".txt", ".xml")):
            txt = open(os.path.join(target, f), encoding="utf-8").read()
            ck.ok("no em dash in %s" % f, "\u2014" not in txt)
            ck.ok("no 'distractor' in %s" % f, "distractor" not in txt.lower())
            for w in ("truncat", "Lynden", "product-limit", "maturity", "follow-up time"):
                ck.ok("no method word %s in %s" % (w, f), w.lower() not in txt.lower())
    # ------------------------------------------------------------ the twelve windows and their bins
    for r in A.RISKS:
        tm = list(C.TARGET_MONTHS[r])
        ck.eq("answer %s" % r, g["windows"][r], tm)
        for p, m in zip((0.8, 0.95), tm):
            d = g["lb"][r][p]
            ck.ok("%s %s cut inside month" % (r, p), 30 * (m - 1) + 1 <= d["q"] <= 30 * m - 6, d)
            ck.ok("%s %s cut clears the calendar-month reading" % (r, p), d["q"] > 30.4375 * (m - 1) + 1, d)
            ck.ok("%s %s share at month start" % (r, p), d["F_lo"] <= p - 0.005, d)
            ck.ok("%s %s share at month end" % (r, p), d["F_hi"] >= p + 0.005, d)
            ck.eq("%s %s interpolated reading" % (r, p), E.month_of(d["qi"]), m)
            ck.record("%s_%s" % (r, p), dict(q=d["q"], F_lo=round(d["F_lo"], 4), F_hi=round(d["F_hi"], 4)))
    # ------------------------------------------------------------ C1: every rigorous estimator, every slip
    G1, G2 = A.ipw_G(pack, "listing"), A.ipw_G(pack, "course")
    v = pack.base()
    tmap = pack.term_map("v4")
    for r in A.RISKS:
        tm = list(C.TARGET_MONTHS[r])
        t = A.tto(pack, r)
        ck.eq("%s lb interp" % r, A.cuts(t, "lb_interp"), tm)
        for d in range(0, 61):
            ck.eq("%s truncation shortened %d days" % (r, d), A.cuts(t, "lb_minus", d=d), tm)
        ck.eq("%s monthly, last complete month" % r, A.cuts(t, "lb_month_floor"), tm)
        ck.eq("%s monthly, month of the lock" % r, A.cuts(t, "lb_month_ceil"), tm)
        rows_r = v[v["reaction_pt"].map(tmap) == r]
        caused = set(rows_r[rows_r["action_taken"].isin(["WITHDRAWN", "INTERRUPTED", "DOSE REDUCED"])]["case_no"])
        ck.eq("%s own window ending at an unrelated course end" % r, A.cuts(t, "lb_window", caused=caused), tm)
        ck.record("%s window ending at any recorded stop" % r, A.cuts(t, "lb_window", caused=caused, all=True))
        ck.eq("%s IPW, start distribution of every case" % r, A.cuts(t, "ipw", G=G1), tm)
        ck.eq("%s IPW, course starts" % r, A.cuts(t, "ipw", G=G2), tm)
        _mature_all(ck, r, t, tm, METHODS if not quick else ["inverted_cdf", "linear"])
        slips = {"duplicates unmerged": dict(dups="keep"), "study cases kept": dict(study=True),
                 "post-lock cases kept": dict(postlock=True), "term list v3": dict(terms="v3"),
                 "partial dates to day 1": dict(partial=1), "partial dates to day 15": dict(partial=15)}
        # a misapplication of the one-reaction rule, not a hygiene slip: asserted on the estimator, recorded on cohorts
        tp = A.tto(pack, r, per_term=True)
        ck.eq("%s every mapped term counted as a reaction" % r, A.cuts(tp, "lb"), tm)
        xs, Ts = tp["x"].to_numpy(float), tp["T"].to_numpy(float)
        off = sum(1 for L in range(int(xs.max()), 914)
                  if [E.month_of(E.np_quantile(xs[Ts >= L], p, "inverted_cdf")) for p in (0.8, 0.95)] != tm)
        ck.record("%s every mapped term counted: mature cutoffs off the answer" % r, off)
        if r == "ild":
            slips["latest delivered version"] = dict(version="latest")
            slips["onset from the first dose ever"] = dict(origin="listing")
            slips["every dosage line a course"] = dict(origin="line")
        for nm, kw in slips.items():
            ts = A.tto(pack, r, **kw)
            ck.eq("%s slip %s" % (r, nm), A.cuts(ts, "lb"), tm)
            ck.eq("%s slip %s, 60 days shorter" % (r, nm), A.cuts(ts, "lb_minus", d=60), tm)
            _mature_all(ck, "%s slip %s" % (r, nm), ts, tm, ["inverted_cdf", "linear"])
    # ------------------------------------------------------------ the ILD ladder
    t = A.tto(pack, "ild")
    ck.eq("ILD rung 1 Weibull", A.cuts(t, "weibull"), RUNG["ild"]["weibull"])
    ck.eq("ILD rung 2 reported cuts", A.cuts(t, "naive"), RUNG["ild"]["naive"])
    t26 = t[t["start"] < pd.Timestamp(2026, 1, 1)]
    ck.eq("ILD rung 3 courses started before 2026", A.cuts(t26, "naive"), RUNG["ild"]["drop2026"])
    t12 = t[t["start"] <= pd.Timestamp(2025, 9, 30)]
    ck.eq("ILD rung 4 courses started a year before the lock", A.cuts(t12, "naive"), RUNG["ild"]["year"])
    t18 = t[t["start"] <= pd.Timestamp(2025, 3, 31)]
    ck.record("ILD courses started at least 18 months before the lock", A.cuts(t18, "naive"))
    t36 = t[t["start"] <= pd.Timestamp(2023, 9, 30)]
    ck.record("ILD courses started at least 36 months before the lock", dict(n=len(t36), cuts=A.cuts(t36, "naive")))
    ck.record("ILD n", dict(cases=len(t), started_2023=int((t["start"] <= pd.Timestamp(2023, 12, 31)).sum()),
                            T913=int((t["T"] >= 913).sum()), longest=int(t["x"].max())))
    if not quick:
        ck.record("ILD truncated Weibull", A.cuts(t, "trunc_weibull"))
        ck.record("ILD truncated lognormal", A.cuts(t, "trunc_lognormal"))
    for r in A.RISKS:
        tr = A.tto(pack, r)
        nv = A.cuts(tr, "naive")
        ck.eq("%s reported-cuts rung" % r, nv, NAIVE[r])
        ck.ok("%s both cells move between rung 2 and the answer" % r, nv[0] != g["windows"][r][0] and nv[1] != g["windows"][r][1])
        ck.record("%s reported cuts in days" % r, [E.np_quantile(tr["x"].to_numpy(float), p, "inverted_cdf") for p in (0.8, 0.95)])
    # lens swap: rung 2 and the answer read the same reactions
    ck.eq("lens swap: same reactions under both estimators", len(A.tto(pack, "ild")), len(t))
    # ------------------------------------------------------------ Ask A ladders and direction discipline
    for r, stops in LADDER.items():
        tm = list(C.TARGET_MONTHS[r])
        for name, est, kw in stops:
            got = A.cuts(A.tto(pack, r, **kw), est)
            for i in (0, 1):
                key = (r, i, est, tuple(sorted(kw.items())), kw.get("version", "lock"))
                if key in KNOWN_SAME:
                    ck.eq("%s stop %s cell %d lands on the answer (recorded)" % (r, name, i), got[i], tm[i])
                else:
                    ck.ok("%s stop %s cell %d off the answer" % (r, name, i), got[i] != tm[i], got)
            ck.record("Ask A %s %s" % (r, name), got)
        for est in ("lb", "naive"):
            for dev in [None] + DEVICES[r]:
                for ver in ("lock", "latest"):
                    if est == "lb" and dev is None and ver == "lock":
                        continue
                    kw = dict(dev or {})
                    kw["version"] = ver
                    got = A.cuts(A.tto(pack, r, **kw), est)
                    for i in (0, 1):
                        key = (r, i, est, tuple(sorted((dev or {}).items())), ver)
                        if key in KNOWN_SAME:
                            continue
                        ck.ok("%s direction discipline %s %s %s cell %d" % (r, est, dev, ver, i), got[i] != tm[i], got)
    # ------------------------------------------------------------ Ask B
    gold = {"py_launch": R10(g["py"]["launch"]), "py_interval": R10(g["py"]["interval"]),
            "py_15": R10(g["py_strength"][15]), "py_30": R10(g["py_strength"][30])}
    for k, val in (("py_launch", g["py"]["launch"]), ("py_interval", g["py"]["interval"]),
                   ("py_15", g["py_strength"][15]), ("py_30", g["py_strength"][30])):
        ck.ok("%s clear of its rounding edge" % k, abs((val % 10) - 5.0) >= 1.5, val)
        ck.ok("%s not a round hundred" % k, gold[k] % 100 != 0, gold[k])
    ck.eq("strength split adds up", gold["py_15"] + gold["py_30"], gold["py_launch"])
    for r in A.RISKS:
        for w in ("launch", "interval"):
            rate = g["rates"][r][w]
            ck.ok("%s %s rate clear of its rounding edge" % (r, w), abs(((rate * 10) % 1) - 0.5) / 10 >= 0.015, rate)
            gold["%s_%s" % (r, w)] = R1(rate)
            ck.eq("%s %s rate on the rounded patient-years" % (r, w),
                  R1(1000.0 * g["counts"][r][w] / R10(g["py"][w])), R1(rate))

    def figs(kw, counts=None):
        counts = counts or g["counts"]
        pl, pi = sup.patient_years(window="launch", **kw), sup.patient_years(window="interval", **kw)
        out = {"py_launch": R10(pl), "py_interval": R10(pi), "py_15": R10(sup.patient_years(window="launch", strength=15, **kw)),
               "py_30": R10(sup.patient_years(window="launch", strength=30, **kw))}
        for r in A.RISKS:
            out["%s_launch" % r] = R1(1000.0 * counts[r]["launch"] / pl)
            out["%s_interval" % r] = R1(1000.0 * counts[r]["interval"] / pi)
        return out

    names = list(SUPPLY_DEV)
    for k in range(1, len(names) + 1):
        for sub in itertools.combinations(names, k):
            kw = {}
            for nm in sub:
                kw.update(SUPPLY_DEV[nm])
            f = figs(kw)
            for key in gold:
                if key == "py_15" and set(sub) == {"B-H5 mg per DDD"}:
                    ck.eq("15 mg patient-years untouched by the DDD reading", f[key], gold[key])
                    continue
                ck.ok("supply subset %s moves %s" % ("+".join(sub), key), f[key] != gold[key], (f[key], gold[key]))
    ladder_b = [("natural read: sell-in only, every pack 28 tablets, gross, to 6 October",
                 dict(channels="sellin", pack="all28", returns="gross", cutoff="export")),
                ("half device: sell-in with pack sizes, returns netted", dict(channels="sellin")),
                ("channel found, unit missed", dict(cases=False)),
                ("over-cleaned: zero-value lines dropped", dict(foc="drop")),
                ("dose arithmetic: 30 mg as two DDDs", dict(dose="ddd"))]
    for nm, kw in ladder_b:
        pl = sup.patient_years(window="launch", **kw)
        ck.ok("Ask B stop %s off" % nm, R10(pl) != gold["py_launch"])
        ck.record("Ask B %s" % nm, dict(py=round(pl, 1), vs=round(100 * (pl / g["py"]["launch"] - 1), 1)))
    for nm, kw in (("H2 cases received after the lock", dict(postlock=True)), ("H3 study cases", dict(study=True)),
                   ("H4 duplicates unmerged", dict(dups="keep")), ("H1 latest delivered version", dict(version="latest"))):
        counts = {r: {w: A.case_counts(pack, r, w, **kw) for w in ("launch", "interval")} for r in A.RISKS}
        f = figs({}, counts)
        for r in A.RISKS:
            for w in ("launch", "interval"):
                key = "%s_%s" % (r, w)
                if nm.startswith("H1") and r == "ild":
                    ck.eq("no ILD reaction added after the lock (%s)" % w, counts[r][w], g["counts"][r][w])
                    continue
                ck.ok("count hazard %s moves %s" % (nm, key), f[key] != gold[key], (f[key], gold[key]))
    counts_v3 = {r: {w: A.case_counts(pack, r, w, terms="v3") for w in ("launch", "interval")} for r in A.RISKS}
    ck.eq("term list v3 counts the same cases", counts_v3, g["counts"])
    # ------------------------------------------------------------ separation of the ILD call's rows
    L = pack.L
    ild_pts = {pt for pt, rr in tmap.items() if rr == "ild"}
    pop = set(t["case_no"])
    lines = pack.lines()
    multi = 0
    for c, on in zip(t["case_no"], t["onset"]):
        ln = lines[c]
        starts = pd.to_datetime(ln["start_date"])
        multi += int((starts <= on).sum() > 1)
    ck.eq("ILD rows with a restart, step-up or pause before onset", multi, 0)
    lab = L[(L["reaction_pt"].isin(ild_pts)) & (L["onset_date"] == "")]
    ck.eq("ILD reactions dated from the laboratory", len(lab), 0)
    il = L[L["reaction_pt"].isin(ild_pts)]
    per = il.groupby(["case_no", "reaction_pt"])["onset_date"].nunique()
    ck.eq("ILD onset dates changed by any version", int((per > 1).sum()), 0)
    first_v = il.groupby("case_no")["version_no"].min()
    added = first_v[first_v > 1]
    ck.eq("ILD reactions added by a follow-up", len(added), 0)
    # ------------------------------------------------------------ clean-data tests
    tmpd = tempfile.mkdtemp()
    try:
        shutil.copytree(target, os.path.join(tmpd, "t"))
        tt = os.path.join(tmpd, "t")
        lst = pd.read_csv(os.path.join(tt, P.FILES["listing"]), dtype=str, keep_default_na=False)
        lst = lst[lst["version_receipt_date"] <= "2026-09-30"]
        P.write_csv(os.path.join(tt, P.FILES["listing"]), lst)
        si = pd.read_csv(os.path.join(tt, P.FILES["sellin"]), dtype=str, keep_default_na=False)
        P.write_csv(os.path.join(tt, P.FILES["sellin"]), si[si["posting_period"] <= "2026-09"])
        for k in P.DISTRACTORS:
            os.remove(os.path.join(tt, P.FILES[k]))
        pk2, sup2, g2 = G.solve(tt)
        for r in A.RISKS:
            ck.eq("clean data: %s answer unchanged" % r, g2["windows"][r], g["windows"][r])
            ck.eq("clean data: %s reported cuts unchanged" % r, A.cuts(A.tto(pk2, r), "naive"), A.cuts(A.tto(pack, r), "naive"))
            ck.eq("clean data: %s counts unchanged" % r, g2["counts"][r], g["counts"][r])
        ck.eq("clean data: launch patient-years unchanged", round(g2["py"]["launch"], 6), round(g["py"]["launch"], 6))
        ck.ok("clean data: the answer still differs from the reported cuts", g2["windows"]["ild"] != A.cuts(A.tto(pk2, "ild"), "naive"))
    finally:
        shutil.rmtree(tmpd)
    # ------------------------------------------------------------ dates in the pack
    allow = {P.FILES["signals"]}
    iso = re.compile(r"(20\d\d-\d\d-\d\d)")
    for f in files:
        if f.endswith((".csv", ".xml", ".txt")):
            mx = max(iso.findall(open(os.path.join(target, f), encoding="utf-8").read()) or ["0"])
            ck.ok("no date after the extraction run in %s" % f, mx <= "2026-10-06" or f in allow, mx)
    house(ck, target, files, pack, sup, g)
    import openpyxl
    for f in files:
        if not f.endswith(".xlsx"):
            continue
        wb = openpyxl.load_workbook(os.path.join(target, f), read_only=True)
        for ws in wb.worksheets:
            hdr = None
            for row in ws.iter_rows(values_only=True):
                if hdr is None and row and any(isinstance(v, str) for v in row):
                    hdr = list(row)
                for j, val in enumerate(row):
                    if isinstance(val, dt.datetime):
                        col = hdr[j] if hdr and j < len(hdr) else None
                        if col == "next_review":
                            continue
                        ck.ok("no date after the extraction run in %s %s" % (f, ws.title), val <= dt.datetime(2026, 10, 6), (col, val))
        wb.close()
    return ck, g


RETIRED_NAMES = ["Nordisk Centrallaboratorium", "Laborgemeinschaft Nordring", "Labor Isartor", "Amstelland", "Sévigné",
                 "Biopôle Atlantique", "Porta Nuova", "Partenope", "Clínico Turia", "Hanse Pharma"]


def _text_of(path):
    if path.endswith((".csv", ".txt", ".xml")):
        return open(path, encoding="utf-8").read()
    if path.endswith((".xlsx", ".docx")):
        import zipfile
        with zipfile.ZipFile(path) as z:
            return "".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml"))
    return ""


def _max_date(path):
    """Latest date a file carries: ISO dates in text files, date cells in workbooks (next_review excepted)."""
    if path.endswith((".csv", ".txt", ".xml")):
        found = re.findall(r"(20\d\d-\d\d-\d\d)", open(path, encoding="utf-8").read())
        return max(found) if found else None
    if path.endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        best = None
        for ws in wb.worksheets:
            hdr = None
            for row in ws.iter_rows(values_only=True):
                if hdr is None and row and any(isinstance(v, str) for v in row):
                    hdr = list(row)
                for j, val in enumerate(row):
                    if isinstance(val, dt.datetime) and not (hdr and j < len(hdr) and hdr[j] == "next_review"):
                        d = val.date().isoformat()
                        best = d if best is None or d > best else best
        wb.close()
        return best
    return None


def house(ck, target, files, pack, sup, g):
    """The reduce-house-fixes register wired in: records behind every supporting figure of the design note (H2), the
    pre-onset gap clearance, the package log as the pack's register (H9), each file's dates against its own run
    (H16) and the retired invented names (H21)."""
    # H9: the data package log registers every other file of the pack, and only those
    log = open(os.path.join(target, P.FILES["log"]), encoding="utf-8").read()
    logged = set(re.findall(r"([A-Za-z0-9_.-]+\.(?:csv|xlsx|pdf|docx|xml))", log))
    ck.eq("H9 the package log names every other file", logged, set(files) - {P.FILES["log"]})
    # H16: no file carries a date after its own extraction run as the log states it
    mon = {"JAN": 1, "FEB": 2, "MAR": 3, "APR": 4, "MAY": 5, "JUN": 6, "JUL": 7, "AUG": 8, "SEP": 9, "OCT": 10, "NOV": 11, "DEC": 12}
    runs = 0
    for m in re.finditer(r"^\d\d\s+.*?(\d\d)-([A-Z]{3})-(\d{4}) \d\d:\d\d\s+[\d,]+\s+(\S+)$", log, re.M):
        run_day = dt.date(int(m.group(3)), mon[m.group(2)], int(m.group(1))).isoformat()
        mx = _max_date(os.path.join(target, m.group(4)))
        ck.ok("H16 %s dated within its run of %s" % (m.group(4), run_day), mx is None or mx <= run_day, mx)
        runs += 1
    ck.eq("H16 every extraction run read from the log", runs, 13)
    # H21: no retired invented name anywhere in the pack
    for f in files:
        txt = _text_of(os.path.join(target, f))
        for nm in RETIRED_NAMES:
            ck.ok("H21 retired name %s absent from %s" % (nm, f), nm not in txt)
    # pre-onset gaps clear the annex's 28 days on both sides, so every gap convention builds the same courses
    lines = pack.lines()
    worst = []
    for r in A.RISKS:
        t = A.tto(pack, r)
        n_restart = n_step = n_pause = 0
        for c, on in zip(t["case_no"], t["onset"]):
            ln = lines[c]
            st = pd.to_datetime(ln["start_date"]).tolist()
            sp = [pd.Timestamp(x) if len(x) == 10 else None for x in ln["stop_date"]]
            ends = ln["line_end"].tolist()
            restart = step = pause = False
            for i in range(1, len(st)):
                if st[i] > on:
                    break
                gap = (st[i] - sp[i - 1]).days - 1
                ck.ok("pre-onset gap clear of the 28-day rule", gap <= 25 or gap >= 34, (r, c, gap))
                worst.append(gap)
                if gap >= 34:
                    restart = True
                elif ends[i - 1] == "DOSE CHANGED":
                    step = True
                else:
                    pause = True
            n_restart += restart
            n_step += step
            n_pause += pause
        x = t["x"].to_numpy(dtype=int)
        n = len(t)
        months = np.bincount((x - 1) // 30, minlength=25).tolist()
        late = int((x >= 85).sum())
        dated = n - len(A.tto(pack, r, lab="drop")) if r in ("hepatic", "neutropenia") else 0
        ck.record("%s reactions in the estimate" % r, dict(n=n, longest=int(x.max()), months=months))
        ck.record("%s pre-onset course structure" % r,
                  dict(restart=n_restart, step_up=n_step, pause=n_pause,
                       shares=[round(100.0 * v / n, 1) for v in (n_restart, n_step, n_pause)]))
        if dated:
            ck.record("%s laboratory-dated from day 85" % r, dict(dated=dated, from_day85=late, share=round(100.0 * dated / late, 1)))
    ck.record("pre-onset gaps: largest under 28 and smallest over it",
              [max(v for v in worst if v <= 28), min(v for v in worst if v > 28)])
    # supporting figures of the ask sheet
    for nm, kw in SUPPLY_DEV.items():
        pl, pi = sup.patient_years(window="launch", **kw), sup.patient_years(window="interval", **kw)
        ck.record("Ask B single device %s" % nm,
                  dict(launch=round(100 * (pl / g["py"]["launch"] - 1), 1), interval=round(100 * (pi / g["py"]["interval"] - 1), 1),
                       rates_launch=round(100 * (g["py"]["launch"] / pl - 1), 1), rates_interval=round(100 * (g["py"]["interval"] / pi - 1), 1)))
    ck.record("rates unrounded and counts", {r: dict(rates=[round(g["rates"][r][w], 4) for w in ("launch", "interval")],
                                                     counts=[g["counts"][r][w] for w in ("launch", "interval")]) for r in A.RISKS})
    L = pack.L
    first_seen = L.groupby(["case_no", "reaction_pt"])["vrec"].min().reset_index()
    case_first = L.groupby("case_no")["irec"].min()
    tmap = pack.term_map("v4")
    late_rx = first_seen[(first_seen["vrec"] > LOCK) & (first_seen["case_no"].map(case_first) <= LOCK)]
    late_rx = late_rx.assign(risk=late_rx["reaction_pt"].map(tmap)).dropna(subset=["risk"])
    ck.record("risk reactions first added by a post-lock follow-up", late_rx.groupby("risk")["case_no"].nunique().to_dict())
    import world as W
    tot = last12 = old24 = 0.0
    for wk in W.weeks():
        for mk in W.MARKETS:
            v = W.weekly_starts(wk, mk)
            tot += v
            last12 += v if wk > dt.date(2025, 9, 30) else 0.0
            old24 += v if wk <= dt.date(2024, 9, 30) else 0.0
    ck.record("course starts from launch (generator)", dict(total=round(tot), last12=round(100 * last12 / tot, 1),
                                                         over24=round(100 * old24 / tot, 1)))


if __name__ == "__main__":
    import sys
    tgt = sys.argv[1] if len(sys.argv) > 1 else os.path.join(G.TASK, "target")
    ck, g = run(tgt, quick="--quick" in sys.argv)
    print("assertions passed:", ck.n)
    print(json.dumps(ck.rec, indent=1, default=str)[:6000])
