"""task235 documents: the metrics catalogue's purchase-tracking entry (pdf), the build register (xlsx), the Tern SDK's
release notes, the data dictionary and the readme (md)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=13, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=14, leading=18, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=14, spaceBefore=8,
                        spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10, textColor=colors.HexColor("#444444"))
    return dict(body=body, h1=h1, h2=h2, small=small)


def _footer(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(18 * mm, 10 * mm, text)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _table(rows, widths, size=8.0):
    t = Table(rows, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", size), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", size),
                           ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF0")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


# ------------------------------------------------------------------ the metrics catalogue entry
def catalogue(path):
    S = _styles(); b, h1, h2, sm = S["body"], S["h1"], S["h2"], S["small"]
    st = [Paragraph("KC-DQ-04 Purchase tracking quality", h1),
          Paragraph("%s metrics catalogue, %s, Analytics. Entry owner: %s, data quality lead. Version 3, in force from "
                    "1 September 2026; replaces version 2 (March 2026)." % (P.GAME, P.STUDIO, P.PEOPLE["dq"]), sm),
          Spacer(1, 4),
          Paragraph("1. What the check is for", h2),
          Paragraph("Experiment readouts, the purchase funnels and the revenue dashboards read purchases from the game's "
                    "analytics SDK, Tern, through its purchase_completed event. The purchase service checks every purchase "
                    "with the App Store or Google Play as it happens, and its validations are the studio's record of what "
                    "players bought. This entry sets how the two are compared for each production build, and when a build "
                    "is alerted.", b),
          Paragraph("2. Purchases", h2),
          Paragraph("A purchase is a validation of transaction kind purchase in the store's Production environment with "
                    "purchase type standard. A Harbour Pass renewal is made by the store on the player's behalf, not in a "
                    "build, and is not a purchase made in a build. Purchases by the studio's testers (App Store sandbox "
                    "purchases and Google Play licence-tester purchases, purchase type test) are not purchases. A refunded "
                    "purchase remains a purchase. A purchase belongs to the build that made it (the validation's app_build) "
                    "and to the Irish date (Europe/Dublin) of its purchased_at.", b),
          Paragraph("3. A purchase's event", h2),
          Paragraph("A purchase's event is the purchase_completed event that carries its store transaction id. The "
                    "warehouse can hold more than one row for it, when the SDK sends a batch again or the game logs a "
                    "purchase again as it restores it on a new install; the rows are one event.", b),
          Paragraph("4. Tracking loss", h2),
          Paragraph("A build's tracking loss over a period is the share of its purchases in the period whose "
                    "purchase_completed event never reaches the warehouse. Purchases made shortly before the run count "
                    "like any other: an event can reach the warehouse well after its purchase, and one that has not "
                    "arrived by the run is lost only if it never will.", b),
          Paragraph("5. The weekly alert", h2),
          Paragraph("Every Monday at 07:00 the data quality lead runs the check over the four weeks to the Sunday before "
                    "(Irish dates) for every build in production in the build register. A build whose tracking loss over "
                    "the four weeks is above 2 per cent, compared unrounded, gets a tracking-loss alert. An alerted build's "
                    "purchase events are left out of experiment readouts and the revenue funnels until a fixed build "
                    "replaces it, the SDK lead opens a P1 bug, and a build in a staged rollout is held at its current "
                    "share. The run posts the alert list, one row for each production build, to the analytics channel.", b),
          Paragraph("6. Reported with the check", h2),
          Paragraph("Each build's paying players over the period are the distinct players with at least one purchase in "
                    "the build in the period. A player is counted once however many accounts they have used: a guest "
                    "account is the same player as the account it signs in to, two accounts merged by player support are "
                    "the same player, and so on through any chain of these. The account service's links record each move of "
                    "an account's history to another account; they are not a register of every account a player has used. "
                    "Each build's gross bookings per paying player "
                    "are its gross bookings over the period divided by its paying players, to the cent. Gross bookings are "
                    "the customer prices of the build's purchases in the period, before refunds, each converted to euro at "
                    "the ECB euro reference rate in force on the purchase's Irish date (the latest rate published on or "
                    "before that date) and rounded to the cent.", b),
          Paragraph("7. Roles", h2),
          _table([["Role", "Holder", "In this entry"],
                  ["Data quality lead", P.PEOPLE["dq"], "runs the check and raises the alerts (s.5)"],
                  ["SDK lead", P.PEOPLE["sdk"], "opens the P1 bug on an alerted build (s.5)"],
                  ["Release manager", P.PEOPLE["release"], "holds a staged rollout at its current share (s.5)"]],
                 [48 * mm, 40 * mm, 80 * mm]),
          ]
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title="KC-DQ-04 Purchase tracking quality", author=P.PEOPLE["dq"],
                            creator=P.PEOPLE["dq"], invariant=1)
    foot = "KC-DQ-04 v3 | %s metrics catalogue | %s" % (P.GAME, P.STUDIO)
    doc.build(st, onFirstPage=_footer(foot), onLaterPages=_footer(foot))
    postprocess.scrub(path, P.PEOPLE["dq"], "2026-08-27")


# ------------------------------------------------------------------ the build register
def register(path):
    wb = Workbook(); ws = wb.active; ws.title = "Builds"
    ws.append(["platform", "version", "build_number", "released", "sdk_version", "minimum_os", "status", "rollout"])
    for b in P.BUILDS:
        roll = "staged, %d per cent from %s" % (int(P.ROLLOUT[-1][1] * 100), P.ROLLOUT[-1][0].strftime("%d %b %Y")) \
            if b[1] == "7.42.0" else "full"
        ws.append([b[0], b[1], b[2], b[3].date(), b[4], b[5], "production", roll])
    for c in ws[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="E8EEF0")
    for i, w in enumerate([10, 9, 13, 12, 12, 12, 12, 30], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=4).number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"
    wr = wb.create_sheet("Rollout 7.42.0")
    wr.append(["from", "share_of_automatic_updates", "note"])
    notes = ["released to 30 per cent of automatic updates", "raised after the first night's crash report",
             "raised; 100 per cent planned for Wednesday 11 November"]
    for (t, s), n in zip(P.ROLLOUT, notes):
        wr.append([t.strftime("%Y-%m-%d %H:%M UTC"), s, n])
    for c in wr[1]:
        c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="E8EEF0")
    for i, w in enumerate([22, 28, 52], 1):
        wr.column_dimensions[get_column_letter(i)].width = w
    wb.save(path)
    postprocess.scrub(path, P.PEOPLE["release"], "2026-11-08")
    postprocess.finish_office(path, "2026-03-02T10:12:00Z", "2026-11-08T09:20:00Z", P.STUDIO)


# ------------------------------------------------------------------ the SDK's release notes
SDK_NOTES = """# Tern SDK changelog

Fogbell Games' in-house analytics SDK. Maintained by the SDK team (Padraic Cormick). Newest first.

## 5.0.1 (28 October 2026)

- iOS: `purchase_completed` is sent as soon as StoreKit reports the transaction, instead of with the next batch.
- Android: uploads now run through WorkManager, at most every 15 minutes while the game is in the foreground and when it
  starts. A batch that has not gone out when the game leaves the foreground is kept with the queue.
- Fixed: `level_complete` was logged twice when a level was replayed from the map on some Android 9 devices.
- Fixed: `screen_view` kept the previous screen's name when the shop was opened from a push notification.

## 5.0.0 (14 October 2026)

- New upload pipeline: batches are compressed, and the warehouse acknowledges each batch by its first and last event id.
- Batches of more than 500 events are split before upload.
- `ad_impression` gains `ad_network` and `placement`.
- Removed: `logRevenue()`, deprecated since 4.4.0.

## 4.6.3 (4 August 2026)

- Each event now carries an id generated on the device (UUID v4). A batch whose acknowledgement does not arrive is sent
  again with the same ids.
- `client_ts` is the device clock in milliseconds since the Unix epoch; the warehouse stamps `received_at` itself.
- Fixed: a crash at start on iOS 15 when the phone's storage was full.

## 4.6.2 (9 June 2026)

- Fixed: `app_version` was empty on the first event after an update on some Android devices.

## 4.6.0 (12 May 2026)

- Events are written to a queue on the device and uploaded in batches when the game starts and every 30 seconds while
  it is in the foreground. The queue is kept when the game is closed or the phone restarts.
- The queue holds up to 5,000 events; past that the oldest are dropped. Uninstalling the game deletes the queue.
- Events carry the account signed in on the phone: the guest account while the game is played as a guest, and the
  player's account from the moment they sign in on it.
- Unity plugin built against Unity 2022.3 LTS.
- Event names over 40 characters are rejected with a console warning.
""" % {"sdk": P.PEOPLE["sdk"]}


