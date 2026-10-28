"""The golden deliverables: free_plan_limit_2027.xlsx (the sizing behind the call, the stopped accounts' year and each
app's recurring revenue bridge) and free_plan_limit_case_2027.docx (the case for the executive)."""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import params as P
import postprocess
import solve as S

XLSX = "free_plan_limit_2027.xlsx"
DOCX = "free_plan_limit_case_2027.docx"
MADE, SAVED = "2026-10-14T08:30:00Z", "2026-10-16T17:05:00Z"


def gbp(x):
    return "GBP %s" % format(int(round(x)), ",")


def thousand(x):
    return "GBP %s,000" % format(int(round(x / 1000.0)), ",")


def per100(x):
    return int(round(100 * x))


def _sheet(ws, header, rows, widths, fmts=None):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE4EE")
        c.alignment = Alignment(vertical="top", wrap_text=True)
    for r in rows:
        ws.append(list(r))
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for col, f in (fmts or {}).items():
        j = header.index(col) + 1
        for k in range(2, len(rows) + 2):
            ws.cell(row=k, column=j).number_format = f
    ws.freeze_panes = "A2"


def workbook(path, fig):
    c, b, r, ask = fig["call"], fig["by_app"], fig["response"], fig["ask"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Recommendation"
    _sheet(ws, ["Item", "Value"], [
        ["App to take the three-project limit first", c["app"]],
        ["What it adds to %s's 2027 revenue, GBP (nearest 50,000)" % c["app"], c["addition_rounded"]],
        ["As built, GBP", c["addition"]],
        ["Next best app", c["runner_up"]],
        ["What it would add to %s's 2027 revenue, GBP" % c["runner_up"], c["runner_up_addition"]],
        ["Free accounts over three active projects, %s (all accounts)" % fig["kpi_month"], b[c["app"]]["over_three"]],
        ["Free accounts expected to reach a fourth active project during 2027", b[c["app"]]["reaching_four_2027"]],
    ], [62, 18], {"Value": "#,##0"})
    ws = wb.create_sheet("By app")
    rows = []
    for a in S.APPS:
        x = b[a]
        rows.append([a, x["over_three"], x["three_or_fewer"], x["reaching_four_2027"], x["plus_price"], x["free_revenue"],
                     round(r["plus_months_gained"], 3), round(r["free_months_lost"], 3), x["addition"]])
    _sheet(ws, ["App", "Free accounts over three active projects, %s" % fig["kpi_month"],
                "Free accounts with three or fewer, %s" % fig["kpi_month"], "Of those, reaching a fourth in 2027",
                "Plus a month, GBP", "Free-plan revenue per active free account a month, GBP",
                "Months on Plus gained in the year after the limit stops an account",
                "Months of free use lost in that year", "2027 revenue the limit adds, GBP"], rows,
           [12, 16, 16, 14, 10, 16, 18, 14, 16],
           {"Plus a month, GBP": "0.00", "Free-plan revenue per active free account a month, GBP": "0.00",
            "2027 revenue the limit adds, GBP": "#,##0"})
    ws.append(["All five apps", sum(b[a]["over_three"] for a in S.APPS), sum(b[a]["three_or_fewer"] for a in S.APPS),
               sum(b[a]["reaching_four_2027"] for a in S.APPS), None, None, None, None, None])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    ws = wb.create_sheet("Stopped accounts")
    rows = []
    for h in range(12):
        rows.append([h + 1, round(r["stopped_plus"][h], 4), round(r["stopped_free"][h], 4),
                     round(r["unstopped_plus"][h], 4), round(r["unstopped_free"][h], 4)])
    _sheet(ws, ["Months after the month an account reached the storage allowance",
                "Stopped accounts on Plus (share)", "Stopped accounts using the free plan (share)",
                "Unstopped accounts over three projects on Plus (share)",
                "Unstopped accounts over three projects using the free plan (share)"], rows, [22, 16, 18, 18, 20],
           {k: "0.0%" for k in ("Stopped accounts on Plus (share)", "Stopped accounts using the free plan (share)",
                                "Unstopped accounts over three projects on Plus (share)",
                                "Unstopped accounts over three projects using the free plan (share)")})
    ws.append([])
    ws.append(["Free accounts reaching the allowance for the first time, Jan 2023 to Sep 2025 (panel)", r["stops"]])
    ws.append(["Of those, the month after: on Plus, using the free plan, gone", r["plus_next_month"], r["free_next_month"],
               r["gone_next_month"]])
    ws = wb.create_sheet("Bridge")
    rows = []
    for a in S.APPS:
        rows.append([a, ask["%s opening" % a], ask["%s new" % a], ask["%s expansion" % a], -ask["%s contraction" % a],
                     -ask["%s churned" % a], ask["%s closing" % a]])
    _sheet(ws, ["App", "Recurring revenue 30 Sep 2025", "New", "Expansion", "Contraction", "Churned",
                "Recurring revenue 30 Sep 2026"], rows, [12, 18, 12, 12, 12, 12, 18],
           {k: "#,##0;-#,##0" for k in ("Recurring revenue 30 Sep 2025", "New", "Expansion", "Contraction", "Churned",
                                        "Recurring revenue 30 Sep 2026")})
    ws.append(["All five apps"] + [sum(row[j] for row in rows) for j in range(1, 7)])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.number_format = "#,##0;-#,##0"
    ws = wb.create_sheet("Notes")
    notes = [
        ["Records to 30 September 2026; prepared 16 October 2026 for the executive team's meeting of 19 November 2026."],
        ["2027 addition: each app's 2027 revenue with the limit less its 2027 revenue without it (the paper, section 3), a "
         "month on Plus at the Plus price and a month of free use at the free-plan revenue per active free account "
         "(section 4)."],
        ["Accounts reached: free accounts over three active projects in the September 2026 count, and the accounts with "
         "three or fewer expected to reach a fourth during 2027 (the panel's year to September 2026)."],
        ["Stopped account: a free account-month within one per cent of the 2,048 MiB storage allowance, the first in the "
         "panel for each account, January 2023 to September 2025, followed for twelve months; set against free accounts "
         "over three projects below the allowance in the same months."],
        ["Notewell was called Leafnote until April 2025; its earlier months carry the old name in the panel and the count."],
        ["Bridge: GBP a month at the 2026/27 plan rates, on FIN-07. A customer is the person or organisation that pays: a "
         "subscription opened again under a new customer record after the platform ended the old one for failed "
         "payments, and staff Plus subscriptions a company moved onto one Team subscription, stay with that customer."],
    ]
    _sheet(ws, ["Note"], notes, [140])
    for row in ws.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    wb.properties.creator = P.PEOPLE["analyst"]
    wb.properties.lastModifiedBy = P.PEOPLE["analyst"]
    wb.save(path)
    postprocess.scrub(path, P.ORG, SAVED[:10])
    postprocess.finish_office(path, MADE, SAVED, P.ORG, person=P.PEOPLE["analyst"])


def case(path, fig):
    c, b, r = fig["call"], fig["by_app"], fig["response"]
    win = c["app"]
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    sec = d.sections[0]
    fp = sec.footer.paragraphs[0]
    fp.text = "Brindlecote Software, Monetisation. Internal: executive team.    Page "
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)
    for line in ("To: Executive team, meeting of 19 November 2026", "From: %s, product analyst, for %s, head of "
                 "monetisation" % (P.PEOPLE["analyst"], P.PEOPLE["head"]), "Date: 16 October 2026",
                 "Subject: The free plan's three-project limit, which app first"):
        q = d.add_paragraph(line)
        q.paragraph_format.space_after = Pt(0)
    h = d.add_paragraph()
    h.paragraph_format.space_before = Pt(10)
    run = h.add_run("%s should take the three-project limit first" % win)
    run.bold = True
    run.font.size = Pt(15)
    d.add_heading("Recommendation", level=2)
    others = [a for a in sorted(S.APPS, key=lambda a: -b[a]["addition"]) if a != win]
    d.add_paragraph(
        "Take the three-project limit first in %s. It adds about %s to %s's 2027 revenue (%s as built), against %s in "
        "%s, %s in %s, %s in %s and %s in %s." % (
            win, gbp(c["addition_rounded"]), win, gbp(c["addition"]),
            thousand(b[others[0]]["addition"]), others[0], thousand(b[others[1]]["addition"]), others[1],
            thousand(b[others[2]]["addition"]), others[2],
            ("a loss of " + thousand(-b[others[3]]["addition"])) if b[others[3]]["addition"] < 0 else thousand(b[others[3]]["addition"]),
            others[3]))
    d.add_heading("How a free account answers a limit", level=2)
    d.add_paragraph(
        "The free plan already stops accounts once, at the two-gigabyte storage allowance: an account that is full "
        "cannot add files until it deletes some or moves to Plus. In the research panel %s free accounts reached the "
        "allowance for the first time between January 2023 and September 2025. The month after, %.1f per cent were "
        "on Plus, %.1f per cent were still using the free plan and %.1f per cent had stopped using the app. Over the following year they spent "
        "%.2f more months on Plus and %.1f fewer months using the free plan than free accounts over three projects the "
        "allowance had not stopped. The project limit stops an account in the same way: over three projects it cannot "
        "edit until it archives, pays or goes. Free accounts that reached the allowance answered it alike whatever the "
        "app, the number of projects they had or how long they had been with us."
        % (format(r["stops"], ","), fig["shares_pct"]["plus_next_month"], fig["shares_pct"]["free_next_month"],
           fig["shares_pct"]["gone_next_month"], r["plus_months_gained"], r["free_months_lost"]))
    d.add_paragraph(
        "Free accounts over three projects move to Plus on their own at about %.1f per cent a year, and few leave. "
        "Priced at that rate, or at the %.1f per cent who pay when stopped, with nobody leaving, the limit belongs in Tracewood, which "
        "has the most free accounts over the limit. The accounts that leave decide it: each costs its app the free-plan "
        "revenue it earned, GBP 2.00 a month in Tracewood and GBP 0.80 in Notewell, against GBP %.2f in %s, where the "
        "Plus price (GBP %.2f) is the highest." % (fig["shares_pct"]["own_plus_12"], fig["shares_pct"]["plus_next_month"],
                                                  b[win]["free_revenue"], win, b[win]["plus_price"]))
    t = d.add_table(rows=1, cols=5)
    t.style = "Light Grid Accent 1"
    hdr = ["App", "Free accounts over three projects, Sep 2026", "Reaching a fourth during 2027", "Plus a month (GBP)",
           "2027 revenue added (GBP)"]
    for j, v in enumerate(hdr):
        t.rows[0].cells[j].text = v
    for a in S.APPS:
        cells = t.add_row().cells
        x = b[a]
        vals = [a, format(x["over_three"], ","), format(x["reaching_four_2027"], ","), "%.2f" % x["plus_price"],
                format(x["addition"], ",")]
        for j, v in enumerate(vals):
            cells[j].text = v
            if j:
                cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    src = d.add_paragraph("Source: research panel, January 2023 to September 2026; free accounts by active projects, "
                          "September 2026; the paper, section 4. Notewell was called Leafnote until April 2025, and its "
                          "earlier months carry the old name.")
    src.runs[0].font.size = Pt(8.5)
    d.add_heading("What the figure counts", level=2)
    d.add_paragraph(
        "The limit reaches the %s free accounts over three active projects in %s and the %s expected to reach a fourth "
        "during 2027, each from the month after the limit stops it to the end of the year. A month on Plus counts at "
        "the Plus price and a month of free use at the free-plan revenue per active free account, both from the "
        "paper. The counts are every account; the year after a stop is the panel's, one account in twenty-five. Read "
        "with every stop rather than the first, with every unstopped free account as the comparison, or with the stops "
        "of any one year, %s stays first and its addition rounds to the same figure." % (
            format(b[win]["over_three"], ","), win, format(b[win]["reaching_four_2027"], ","), win))
    d.add_heading("Recurring revenue", level=2)
    d.add_paragraph("Each app's recurring revenue bridge for the year to 30 September 2026, on FIN-07, is in "
                    "free_plan_limit_2027.xlsx for the board pack.")
    d.core_properties.author = P.PEOPLE["analyst"]
    d.core_properties.title = "Free plan limit: which app first"
    d.save(path)
    postprocess.scrub(path, P.ORG, SAVED[:10])
    postprocess.finish_office(path, MADE, SAVED, P.ORG, person=P.PEOPLE["analyst"])


def write(folder, fig):
    workbook(os.path.join(folder, XLSX), fig)
    case(os.path.join(folder, DOCX), fig)
