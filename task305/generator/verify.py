"""The independent verifier: reads target.zip (no pandas, nothing from solve.py), opens every file in its own format and
recomputes the call's addition for each app and every bridge figure, against figures.json.

    python3 task305/generator/verify.py [target.zip]
"""
import csv
import datetime as dt
import email
import io
import json
import os
import subprocess
import sys
import tempfile
import zipfile

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
APPS = ["Tracewood", "Taskfold", "Ledgerline", "Formbox", "Notewell"]


def rows(z, name):
    return list(csv.DictReader(io.TextIOWrapper(z.open(name), encoding="utf-8")))


def main(zp):
    fig = json.load(open(os.path.join(HERE, "figures.json")))
    bad = []

    def ok(cond, msg):
        print(("ok    " if cond else "FAIL  ") + msg)
        if not cond:
            bad.append(msg)
    z = zipfile.ZipFile(zp)
    names = sorted(z.namelist())
    ok(len(names) >= 10, "%d files in the zip" % len(names))
    tmp = tempfile.mkdtemp()
    for n in names:
        if n.endswith(".csv"):
            r = rows(z, n)
            ok(len(r) > 0, "%s: %d rows" % (n, len(r)))
        elif n.endswith(".xlsx"):
            wb = load_workbook(io.BytesIO(z.read(n)), read_only=True)
            ok(len(wb.sheetnames) >= 1, "%s: sheets %s" % (n, wb.sheetnames))
        elif n.endswith(".pdf"):
            p = os.path.join(tmp, os.path.basename(n))
            open(p, "wb").write(z.read(n))
            txt = subprocess.run(["pdftotext", "-layout", p, "-"], capture_output=True, text=True).stdout
            ok(len(txt) > 400, "%s: %d characters of text" % (n, len(txt)))
        elif n.endswith(".eml"):
            msg = email.message_from_bytes(z.read(n))
            ok(bool(msg["Subject"]) and bool(msg.get_payload()), "%s: %s" % (n, msg["Subject"]))
    # the paper's table
    p = os.path.join(tmp, "free_plan_limit_paper_2027.pdf")
    txt = subprocess.run(["pdftotext", "-layout", p, "-"], capture_output=True, text=True).stdout
    price, free = {}, {}
    for line in txt.splitlines():
        w = line.split()
        if w and w[0] in APPS and len(w) > 2:
            try:
                price[w[0]], free[w[0]] = float(w[1]), float(w[2])
            except ValueError:
                pass
    # the count of free accounts over three projects and with three or fewer, latest month
    wb = load_workbook(io.BytesIO(z.read("free_accounts_by_active_projects_2023_2026.xlsx")), read_only=True)
    ws = wb[sorted(wb.sheetnames)[-1]]
    data = list(ws.iter_rows(values_only=True))
    last = max(r[0] for r in data[1:])
    heavy = {a: 0 for a in APPS}
    light = {a: 0 for a in APPS}
    for r in data[1:]:
        if r[0] != last:
            continue
        a = "Notewell" if r[1] == "Leafnote" else r[1]
        (light if str(r[2]) in ("1", "2", "3") else heavy)[a] += int(r[3])
    # the panel, account by account
    acc = {}
    months = set()
    for r in rows(z, "research_panel_account_months_2023_2026.csv"):
        y, m = r["month"].split("-")
        mi = int(y) * 12 + int(m) - 1
        months.add(mi)
        acc.setdefault(r["account_id"], {})[mi] = (r["plan"], int(r["active_projects"]), float(r["storage_mib"]))
    lo, hi = min(months), max(months)
    end = hi - 12

    def state(rec, mi):
        v = rec.get(mi)
        return None if v is None else (1 if v[0] == "plus" else 0)
    stop_rows, ctrl_rows = [], []
    for a_, rec in acc.items():
        firsts = [mi for mi in sorted(rec) if rec[mi][0] == "free" and rec[mi][2] >= 0.99 * 2048]
        if firsts and firsts[0] <= end:
            stop_rows.append((a_, firsts[0]))
        for mi, (pl, k, s) in rec.items():
            if mi <= end and pl == "free" and k >= 4 and s < 0.99 * 2048:
                ctrl_rows.append((a_, mi))

    def curve(rs):
        P = [0.0] * 12
        F = [0.0] * 12
        for a_, mi in rs:
            for h in range(12):
                s = state(acc[a_], mi + h + 1)
                P[h] += s == 1
                F[h] += s == 0
        return [x / len(rs) for x in P], [x / len(rs) for x in F]
    P1, F1 = curve(stop_rows)
    P0, F0 = curve(ctrl_rows)
    ok(len(stop_rows) == fig["response"]["stops"], "first stops %d" % len(stop_rows))
    dP = [a - b for a, b in zip(P1, P0)]
    dF = [b - a for a, b in zip(F1, F0)]
    # accounts below four in September a year before the extract, first reaching four month by month
    j0 = end
    base = [a_ for a_, rec in acc.items() if rec.get(j0, ("x", 9, 0))[0] == "free" and rec[j0][1] <= 3]
    fr = [0] * 12
    for a_ in base:
        for h in range(1, 13):
            v = acc[a_].get(j0 + h)
            if v and v[0] == "free" and v[1] >= 4:
                fr[h - 1] += 1
                break
    fr = [x / len(base) for x in fr]
    for a in APPS:
        val = heavy[a] * (sum(dP) * price[a] - sum(dF) * free[a])
        for h in range(1, 12):
            val += light[a] * fr[h - 1] * (sum(dP[:12 - h]) * price[a] - sum(dF[:12 - h]) * free[a])
        ok(abs(val - fig["by_app"][a]["addition"]) <= 1.0, "%s addition %.0f against %d" % (a, val, fig["by_app"][a]["addition"]))
    # the bridge
    rates = {"GBP": 1.0}
    for r in rows(z, "fx_plan_rates.csv"):
        if r["plan_year"] == "2026/27":
            rates[r["currency"]] = float(r["units_per_gbp"])

    def val(r, started):
        v = float(r["unit_amount"]) * int(r["quantity"]) / (12.0 if r["billing_interval"] == "year" else 1.0)
        if r["discount_percent"]:
            e = dt.datetime.strptime(r["discount_ends"], "%Y-%m-%d")
            if (e.year - started.year) * 12 + e.month - started.month + (e.day - started.day) / 31.0 >= 11.99:
                v *= 1 - float(r["discount_percent"]) / 100.0
        return v / rates[r["currency"]]
    ts = lambda s: dt.datetime.strptime(s, "%Y-%m-%dT%H:%M:%SZ")
    cur, start, who = {}, {}, {}
    opening = {}
    for r in rows(z, "billing_subscriptions_20250930.csv"):
        start[r["subscription_id"]] = ts(r["started_at"])
        cur[r["subscription_id"]] = r
        who[r["subscription_id"]] = (r["app"], r["customer_id"], r["account_id"])
        opening[r["subscription_id"]] = val(r, start[r["subscription_id"]])
    live = set(cur)
    for r in rows(z, "billing_subscription_events_20251001_20260930.csv"):
        sid, kind = r["subscription_id"], r["event_type"].split(".")[1]
        who.setdefault(sid, (r["app"], r["customer_id"], r["account_id"]))
        if kind == "created":
            start[sid] = ts(r["occurred_at"])
            live.add(sid)
        if kind in ("created", "updated", "renewed"):
            cur[sid] = r
        if kind == "ended":
            live.discard(sid)
    closing = {sid: val(cur[sid], start[sid]) for sid in live}
    parent = {}

    def find(x):
        while parent.get(x, x) != x:
            x = parent[x]
        return x
    for sid, (app, c, a_) in who.items():
        ra, rb = find(("c", app, c)), find(("a", app, a_))
        if ra != rb:
            parent[ra] = rb
    tot = {}
    for sid, v in opening.items():
        k = (who[sid][0], find(("c", who[sid][0], who[sid][1])))
        tot.setdefault(k, [0.0, 0.0])[0] += v
    for sid, v in closing.items():
        k = (who[sid][0], find(("c", who[sid][0], who[sid][1])))
        tot.setdefault(k, [0.0, 0.0])[1] += v
    for a in APPS:
        n = e = c_ = ch = 0.0
        for (app, _), (o, cl) in tot.items():
            if app != a:
                continue
            if o <= 1e-9 < cl:
                n += cl
            elif cl <= 1e-9 < o:
                ch += o
            elif cl > o:
                e += cl - o
            else:
                c_ += o - cl
        for k, v in (("new", n), ("expansion", e), ("contraction", c_), ("churned", ch)):
            ok(abs(round(v) - fig["ask"]["%s %s" % (a, k)]) <= 0.5, "%s %s %d" % (a, k, round(v)))
    print("%d failures" % len(bad))
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(HERE), "target.zip")))
