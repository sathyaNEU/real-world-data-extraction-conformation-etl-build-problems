"""Cut the task305 pack: the paper and the thread, the price list, the monthly count of free accounts, the research
panel, Billing's subscription records, Finance's policy and plan rates, and two files the organisation keeps for other
work (app-store ratings, acquisition spend)."""
import csv
import datetime as dt
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import billing as BL
import docs
import params as P
import postprocess
import world as W

FILES = P.FILES
DISTRACTORS = ("ratings", "acquisition")
UTC = dt.timezone.utc


def _csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)


def _xlsx(path, sheets, made, saved, author):
    wb = Workbook()
    for i, (title, header, rows, widths, fmts) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDE4EE")
            c.alignment = Alignment(vertical="top", wrap_text=True)
        for r in rows:
            ws.append(list(r))
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        for col, f in (fmts or {}).items():
            j = header.index(col) + 1
            for k in range(2, len(rows) + 2):
                ws.cell(row=k, column=j).number_format = f
        ws.freeze_panes = "A2"
    wb.properties.creator = author
    wb.properties.lastModifiedBy = author
    wb.save(path)
    postprocess.scrub(path, P.ORG, saved[:10])
    postprocess.finish_office(path, made, saved, P.ORG)


def u1(x, slot):
    return BL.u(x, 900 + slot)


def app_name(i, m):
    a = P.APPS[i]
    return P.OLD_NAME.get(a, a) if m < P.RENAME_MONTH else a


# --- the research panel -------------------------------------------------------------------------------------------

COUNTRIES = [("GB", 0.64), ("US", 0.11), ("DE", 0.07), ("FR", 0.05), ("IE", 0.04), ("NL", 0.04), ("ES", 0.02),
             ("SE", 0.01), ("AU", 0.01), ("CA", 0.01)]
SOURCES = [("web", 0.52), ("ios", 0.18), ("android", 0.14), ("google", 0.11), ("invite", 0.05)]


def pick(table, v):
    c = 0.0
    for k, p in table:
        c += p
        if v < c:
            return k
    return table[-1][0]


def panel(world, target):
    rows = world["panel_rows"]
    first, last, app_of = {}, {}, {}
    out = []
    for (m, ids, ap, st, k, s, days, since) in rows:
        for j in range(len(ids)):
            x = int(ids[j])
            first.setdefault(x, m)
            last[x] = m
            app_of[x] = int(ap[j])
            out.append((m, x, int(ap[j]), int(st[j]), int(k[j]), float(s[j]), int(days[j])))
    out.sort(key=lambda r: (r[0], r[1]))
    _csv(os.path.join(target, FILES["panel_months"]),
         ["month", "account_id", "app", "plan", "active_projects", "storage_mib", "days_active"],
         [[P.ym(m), BL.acct(x), app_name(a, m), "plus" if st == 1 else "free", k, "%.1f" % s, d]
          for (m, x, a, st, k, s, d) in out])
    # the account register: sign-up, closure and merges
    left = world.get("panel_left", {})
    active_by = {}
    for (m, x, a, st, k, s, d) in out:
        active_by.setdefault((m, a), []).append(x)
    reg = []
    for x in sorted(first):
        app, born = world["panel_acc"][x]
        created = BL.month_start(born) + dt.timedelta(days=int(u1(x, 1) * 27), hours=int(u1(x, 2) * 24),
                                                     minutes=int(u1(x, 3) * 60), seconds=int(u1(x, 4) * 60))
        status, closed, merged = "open", "", ""
        if last[x] < P.M_EXTRACT:
            lm, own = left.get(x, (last[x] + 1, False))
            if own and u1(x, 5) < 0.08:
                pool = [y for y in active_by.get((min(lm, P.M_EXTRACT), app), []) if y != x and first[y] <= lm]
                if pool:
                    merged = BL.acct(pool[int(u1(x, 6) * len(pool))])
                    status = "closed"
                    closed = (BL.month_start(lm) + dt.timedelta(days=int(u1(x, 7) * 27), hours=10)).strftime("%Y-%m-%dT%H:%M:%SZ")
            if not merged and u1(x, 8) < 0.3:
                cm = last[x] + 1 + int(u1(x, 9) * 6)
                if cm <= P.M_EXTRACT:
                    status = "closed"
                    closed = (BL.month_start(cm) + dt.timedelta(days=int(u1(x, 10) * 27), hours=int(u1(x, 11) * 24),
                                                                minutes=int(u1(x, 12) * 60))).strftime("%Y-%m-%dT%H:%M:%SZ")
        reg.append([BL.acct(x), app_name(app, born), created.strftime("%Y-%m-%dT%H:%M:%SZ"), pick(COUNTRIES, u1(x, 13)),
                    pick(SOURCES, u1(x, 14)), status, closed, merged])
    _csv(os.path.join(target, FILES["panel_accounts"]),
         ["account_id", "app_at_signup", "created_utc", "country", "signup_source", "status", "closed_utc", "merged_into"],
         reg)


# --- the monthly count of free accounts by active projects ---------------------------------------------------------

def kpi(world, target):
    old_hdr = ["Month", "Product", "Active projects", "Active free accounts"]
    new_hdr = ["month", "app", "active_projects", "free_accounts_active"]
    sheets = []
    for year in (2023, 2024, 2025, 2026):
        rows = []
        for mm in range(12):
            m = P.month_of(year, mm + 1)
            if m > P.M_EXTRACT:
                break
            c = world["kpi"][m]
            for i in range(5):
                if year <= 2024:
                    label = dt.date(year, mm + 1, 1).strftime("%b-%y")
                    for band, v in (("1", c[i, 1]), ("2", c[i, 2]), ("3", c[i, 3]), ("4+", c[i, 4:].sum())):
                        rows.append([label, app_name(i, m), band, int(v)])
                else:
                    for band, v in (("1", c[i, 1]), ("2", c[i, 2]), ("3", c[i, 3]), ("4-5", c[i, 4:6].sum()),
                                    ("6-9", c[i, 6:10].sum()), ("10+", c[i, 10])):
                        rows.append([P.ym(m), app_name(i, m), band, int(v)])
        hdr = old_hdr if year <= 2024 else new_hdr
        sheets.append((str(year), hdr, rows, [12, 12, 16, 20], {}))
    _xlsx(os.path.join(target, FILES["kpi"]), sheets, "2023-02-03T08:12:00Z", "2026-10-05T07:40:00Z",
          "Product Analytics")


# --- the price list --------------------------------------------------------------------------------------------

def catalogue(target):
    rows = []
    for i, a in enumerate(P.APPS):
        c = P.CODE[a]
        rows.append([a, "%s-FREE" % c, "Free", "", "0.00", "0.00", "0.00", 2, "", 2, 30, 100])
        for plan, iv in (("plus", "month"), ("plus", "year"), ("team", "month"), ("team", "year")):
            g = BL.price(i, "GBP", plan, iv, BL.RISE)
            e = BL.price(i, "EUR", plan, iv, BL.RISE)
            d = BL.price(i, "USD", plan, iv, BL.RISE)
            rows.append([a, BL.plan_code(i, plan, iv), "Plus" if plan == "plus" else "Team (a seat)",
                         "monthly" if iv == "month" else "annual", "%.2f" % g, "%.2f" % e, "%.2f" % d,
                         100 if plan == "plus" else 1000, "", 10 if plan == "plus" else "", 365, 2000])
    _csv(os.path.join(target, FILES["catalogue"]),
         ["app", "plan_code", "plan", "billing", "price_gbp", "price_eur", "price_usd", "storage_gb",
          "active_projects_limit", "collaborators_per_project", "version_history_days", "max_upload_mb"], rows)


# --- Billing ---------------------------------------------------------------------------------------------------

def billing(world, target):
    opening, ev, companies = BL.build(world)

    def code(r):
        return BL.plan_code(r["app"], r["plan"], r["interval"])
    _csv(os.path.join(target, FILES["subs"]),
         ["subscription_id", "customer_id", "account_id", "app", "plan_code", "billing_interval", "currency",
          "unit_amount", "quantity", "discount_percent", "discount_ends", "started_at", "status"],
         [[r["sid"], r["cus"], r["acc"], P.APPS[r["app"]], code(r), r["interval"], r["cur"], "%.2f" % r["unit"],
           r["seats"], r["disc"] or "", r["disc_end"].strftime("%Y-%m-%d") if r.get("disc_end") else "",
           BL.iso(r["start"]), r["status"]] for r in opening])
    _csv(os.path.join(target, FILES["events"]),
         ["event_id", "occurred_at", "event_type", "subscription_id", "customer_id", "account_id", "app", "plan_code",
          "billing_interval", "currency", "unit_amount", "quantity", "discount_percent", "discount_ends", "status"],
         [["evt_" + BL.b36(BL.hid(k + 1, 99) ^ BL.hid(int(r["t"].timestamp()), 98), 16), BL.iso(r["t"]),
           "subscription." + r["kind"], r["sid"], r["cus"], r["acc"], P.APPS[r["app"]], code(r), r["interval"], r["cur"],
           "%.2f" % r["unit"], r["seats"], r["disc"] or "", r["disc_end"].strftime("%Y-%m-%d") if r.get("disc_end") else "",
           r["status"]] for k, r in enumerate(ev)])
    _csv(os.path.join(target, FILES["fx"]), ["plan_year", "currency", "units_per_gbp"],
         [[y, c, "%.4f" % v] for (y, c, v) in BL.FX])
    return opening, ev, companies


# --- files kept for other work -------------------------------------------------------------------------------------

def ratings(target):
    rows = []
    base = {"Tracewood": 4.5, "Taskfold": 4.3, "Ledgerline": 4.6, "Formbox": 4.2, "Notewell": 4.4}
    for m in range(P.M_PANEL0, P.M_EXTRACT + 1):
        for i, a in enumerate(P.APPS):
            for j, store in enumerate(("App Store", "Google Play")):
                key = m * 10 + i * 2 + j
                r = base[a] - 0.15 * j + 0.12 * (BL.u(key, 31) - 0.5) - (0.08 if (a == "Taskfold" and 300 <= m <= 304) else 0)
                n = int(300 + 2400 * BL.u(key, 32) * (1.6 if a == "Tracewood" else 1.0))
                rows.append([P.ym(m), app_name(i, m), store, "%.2f" % r, n, int(n * (0.18 + 0.1 * BL.u(key, 33)))])
    _csv(os.path.join(target, FILES["ratings"]), ["month", "app", "store", "average_rating", "ratings", "reviews"], rows)


def acquisition(target):
    rows = []
    chans = [("Paid search", 1.0), ("Paid social", 0.7), ("Affiliates", 0.35), ("App store ads", 0.45)]
    for mm in range(1, 10):
        m = P.month_of(2026, mm)
        for i, a in enumerate(P.APPS):
            for j, (ch, w) in enumerate(chans):
                key = m * 20 + i * 4 + j
                spend = round(W.signups(i, m) * 2.2 * w * (0.8 + 0.4 * BL.u(key, 41)), -1)
                su = int(spend / (3.2 + 2.0 * BL.u(key, 42) + 1.5 * (a == "Ledgerline")))
                rows.append([P.ym(m), a, ch, spend, su, round(spend / max(su, 1), 2)])
    _xlsx(os.path.join(target, FILES["acquisition"]),
          [("2026", ["month", "app", "channel", "spend_gbp", "sign_ups", "cost_per_sign_up_gbp"], rows,
            [10, 12, 16, 12, 10, 18], {"spend_gbp": "#,##0", "cost_per_sign_up_gbp": "0.00"})],
          "2026-02-04T09:00:00Z", "2026-10-06T16:20:00Z", "Growth Marketing")


def build(target, world=None):
    world = world or W.simulate()
    docs.paper(os.path.join(target, FILES["paper"]))
    postprocess.scrub(os.path.join(target, FILES["paper"]), P.ORG, "2026-10-12")
    docs.policy(os.path.join(target, FILES["policy"]))
    postprocess.scrub(os.path.join(target, FILES["policy"]), P.ORG, "2026-04-14")
    docs.thread(os.path.join(target, FILES["thread"]))
    catalogue(target)
    kpi(world, target)
    panel(world, target)
    billing(world, target)
    ratings(target)
    acquisition(target)
    return world
