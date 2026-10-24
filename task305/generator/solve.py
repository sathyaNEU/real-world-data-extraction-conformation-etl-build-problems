"""Every graded figure from the pack alone: the call (the resident route and the golden, which carries the holiday
parks as origins), and the three asks. Reads only the files in target/."""
import csv
import datetime as dt
import io
import math
import os
import zipfile
import zoneinfo

import numpy as np
from openpyxl import load_workbook

import params as P

LONDON = zoneinfo.ZoneInfo("Europe/London")
UTC = dt.timezone.utc
OFF = ("2025-11", "2025-12", "2026-01", "2026-02")


def rd(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def sheet(path, name):
    ws = load_workbook(path, read_only=True, data_only=True)[name]
    it = ws.iter_rows(values_only=True)
    head = next(it)
    return [dict(zip(head, r)) for r in it if any(v is not None for v in r)]


class Pack:
    def __init__(self, target):
        F = {k: os.path.join(target, v) for k, v in P.FILES.items()}
        self.F = F
        hh = {r["geography code"]: int(r["Number of households"]) for r in rd(F["households"])}
        pwc = {r["OA21CD"]: (float(r["x"]), float(r["y"])) for r in rd(F["pwc"])}
        spend = {r["OA21CD"]: float(r["AvgWeeklyGrocerySpendPerHousehold_GBP_2026"]) for r in rd(F["spend"])}
        self.oas = [(pwc[k], hh[k] * spend[k] * 52 * P.INSTORE) for k in sorted(hh)]
        self.pc = {r["pcds"]: (float(r["oseast1m"]), float(r["osnrth1m"])) for r in rd(F["onspd"])}
        with open(F["voa_list"], encoding="utf-8", newline="") as fh:
            self.voa = [ln.rstrip("\r\n").split("*") for ln in fh if ln.strip()]
        self.smv = {r["uarn"]: (float(r["main_measure"]), r["unit_of_measurement"]) for r in rd(F["voa_smv"])}
        self.estate = sheet(F["estate"], "Stores")
        self.sites = sheet(F["shortlist"], "Shortlist")
        own_pc = {r["postcode"]: str(r["store_no"]) for r in self.estate}
        self.stores = []
        self.parks = []
        for r in self.voa:
            desc, uarn, pc = r[5], r[6], r[14]
            if desc in P.GROCERY and uarn in self.smv and self.smv[uarn][1] == "m2" and self.smv[uarn][0] >= P.MIN_AREA:
                self.stores.append(dict(id=own_pc.get(pc, "u" + uarn), own=pc in own_pc, xy=self.pc[pc], a=self.smv[uarn][0]))
            if uarn in self.smv and self.smv[uarn][1] == "pitches":
                self.parks.append(dict(uarn=uarn, desc=desc, pc=pc, xy=self.pc[pc], n=self.smv[uarn][0]))
        self.cards = rd(F["cards"])

    # ------------------------------------------------------------- the spend-share model
    @staticmethod
    def shares(xy, S):
        X = np.array([s["xy"][0] for s in S]); Y = np.array([s["xy"][1] for s in S])
        A = np.array([s["a"] for s in S], float) ** P.ALPHA
        w = A * np.exp(-P.BETA * P.ROADF * np.hypot(X - xy[0], Y - xy[1]) / 1000.0)
        return w / w.sum()

    def sales(self, S, origins):
        out = np.zeros(len(S))
        for xy, d in origins:
            out += d * self.shares(xy, S)
        return out

    def guests(self, off=OFF, gross=True, stores=None):
        """the guests' in-store spend at each of our stores: out-of-region card spend in the season above the store's
        own November to February level, grossed up by the store's card share"""
        share = {str(r["store_no"]): float(r["card_share_of_sales"]) for r in self.estate}
        nl = {}
        for r in self.cards:
            if not r["cardholder_home_sector"].startswith(P.AREA):
                nl.setdefault(r["store_no"], {}).setdefault(r["month"], 0.0)
                nl[r["store_no"]][r["month"]] += float(r["spend_gbp"])
        out = {}
        for s, mm in nl.items():
            level = np.mean([mm.get(m, 0.0) for m in off]) if off else 0.0
            season = [m for m in P.MONTHS if int(m[5:]) in P.SEASON]
            g = sum(mm.get(m, 0.0) - level for m in season)
            out[s] = g / (share[s] if gross else 1.0)
        return out

    def v_per_pitch(self, guests=None, parks=None, only=None):
        guests = guests if guests is not None else self.guests()
        parks = parks if parks is not None else self.parks
        own = [k for k, s in enumerate(self.stores) if s["own"] and (only is None or s["id"] in only)]
        den = 0.0
        for p in parks:
            sh = self.shares(p["xy"], self.stores)
            den += p["n"] * sum(sh[k] for k in own)
        num = sum(v for s, v in guests.items() if only is None or s in only)
        return num / den

    def appraise(self, V=None, parks=None, margin=P.MARGIN):
        """each site's addition: 26 per cent of its sales less 26 per cent of what it takes from our ten stores, less
        its running costs; V None is the resident route"""
        parks = parks if parks is not None else self.parks
        origins = list(self.oas) + ([(p["xy"], p["n"] * V) for p in parks] if V else [])
        S = self.stores
        own = [k for k, s in enumerate(S) if s["own"]]
        base = self.sales(S, origins)
        out = {}
        for r in self.sites:
            S2 = S + [dict(id=r["site"], own=True, xy=self.pc[r["postcode"]], a=float(r["proposed_floor_area_m2"]))]
            b2 = self.sales(S2, origins)
            took = sum(base[k] - b2[k] for k in own)
            cost = sum(float(r[c]) for c in ("rent_gbp", "business_rates_gbp", "staffing_gbp", "other_running_gbp"))
            out[r["site"]] = dict(new=b2[-1], took=took, cost=cost, add=margin * (b2[-1] - took) - cost)
        return out

    def call(self):
        V = self.v_per_pitch()
        g = self.appraise(V)
        r0 = self.appraise(None)
        win = max(g, key=lambda k: g[k]["add"])
        return dict(V=V, golden=g, resident=r0, site=win, add=g[win]["add"],
                    resident_site=max(r0, key=lambda k: r0[k]["add"]))


# ----------------------------------------------------------------- ask A
def to_local(s):
    t = dt.datetime.strptime(s, "%Y-%m-%d %H:%M").replace(tzinfo=UTC)
    return t.astimezone(LONDON).replace(tzinfo=None)


def _utc(s):
    return dt.datetime.strptime(s, "%Y-%m-%d %H:%M")


def turned_away(rows):
    """the deliveries a store turned away: the stop closed sooner than its cases could have come off the trailer
    (under 4 minutes plus 0.02 of a minute a case), the same cases going to the store again the next day"""
    out = set()
    for r in rows:
        if r["stop_type"] == "DEL":
            dwell = (_utc(r["depart_utc"]) - _utc(r["arrive_utc"])).total_seconds() / 60.0
            if dwell < 4 + 0.02 * int(r["cases"]):
                out.add((r["trip_id"], r["stop_seq"]))
    return out


def ask_a(target, mode="golden"):
    """the cases the DC delivered to each store a week over the 25 weeks, by the stop's UK arrival date: a load a store
    turned away is delivered once, by the next day's trip"""
    rows = rd(os.path.join(target, P.FILES["tms"]))
    dels = [r for r in rows if r["stop_type"] == "DEL" and r["status"] == "COMPLETE"]
    drop = set()
    if mode in ("golden", "utc"):
        drop = turned_away(dels)
    elif mode == "drop_short":
        drop = {(r["trip_id"], r["stop_seq"]) for r in dels
                if (_utc(r["depart_utc"]) - _utc(r["arrive_utc"])).total_seconds() / 60.0 < 15}
    elif mode == "dedup_equal":
        seen = {}
        for r in sorted(dels, key=lambda r: r["arrive_utc"]):
            d = to_local(r["arrive_utc"]).date()
            k = (r["store_no"], r["cases"])
            if k in seen and (d - seen[k]).days == 1:
                drop.add((r["trip_id"], r["stop_seq"]))
            seen[k] = d
    tot = {}
    for r in dels:
        if (r["trip_id"], r["stop_seq"]) in drop:
            continue
        d = to_local(r["arrive_utc"]).date() if mode != "utc" else _utc(r["arrive_utc"]).date()
        if P.TMS_FROM <= d <= P.TMS_TO:
            tot[r["store_no"]] = tot.get(r["store_no"], 0) + int(r["cases"])
    weeks = ((P.TMS_TO - P.TMS_FROM).days + 1) / 7.0
    return {k: v / weeks for k, v in tot.items()}


# ----------------------------------------------------------------- ask B
def _t(date, hhmm):
    return dt.datetime.strptime(date + " " + hhmm, "%Y-%m-%d %H:%M")


def wfm_shifts(target):
    """the workforce export's rows in file order, each with its start and end as datetimes (a shift ending at or
    before its start ends the next day)"""
    out = []
    for i, r in enumerate(rd(os.path.join(target, P.FILES["wfm"]))):
        a, b = _t(r["shift_date"], r["shift_start"]), _t(r["shift_date"], r["shift_end"])
        if b <= a:
            b += dt.timedelta(days=1)
        out.append(dict(r, i=i, a=a, b=b, paid=float(r["paid_hours"])))
    return out


def punches_by_employee(target):
    reg = {r["terminal_id"]: r["store_no"] for r in rd(os.path.join(target, P.FILES["terminals"]))}
    out = {}
    for p in rd(os.path.join(target, P.FILES["punches"])):
        out.setdefault(p["employee_no"], []).append(
            (dt.datetime.strptime(p["punch_time"], "%Y-%m-%d %H:%M"), p["punch"], reg[p["terminal_id"]]))
    return out


def current_shifts(rows, pun, by="punches"):
    """one row per shift worked: rows of one colleague and day whose times overlap are one shift extracted twice,
    and the one standing is the row whose end the clock-out follows (by="punches") or the later row in the file
    (by="file"); rows that do not overlap are separate shifts"""
    groups = {}
    for r in rows:
        groups.setdefault((r["employee_no"], r["shift_date"]), []).append(r)
    keep = []
    for g in groups.values():
        g.sort(key=lambda r: r["i"])
        clusters = []
        for r in g:
            for c in clusters:
                if any(r["a"] < q["b"] and q["a"] < r["b"] for q in c):
                    c.append(r)
                    break
            else:
                clusters.append([r])
        for c in clusters:
            if len(c) == 1 or by == "file":
                keep.append(c[-1])
                continue
            outs = [t for t, kind, _ in pun.get(c[0]["employee_no"], []) if kind == "OUT"]
            ok_ = [r for r in c if any(0 <= (t - r["b"]).total_seconds() / 60 <= 15 for t in outs)]
            keep.append(ok_[-1] if ok_ else c[-1])
    return keep


def worked_store(r, pun):
    """the store of the terminal the colleague clocked in at for the shift, or clocked out at where no clock-in"""
    ps = pun.get(r["employee_no"], [])
    ins = [s for t, kind, s in ps if kind == "IN" and 0 <= (r["a"] - t).total_seconds() / 60 <= 15]
    if ins:
        return ins[0]
    outs = [s for t, kind, s in ps if kind == "OUT" and 0 <= (t - r["b"]).total_seconds() / 60 <= 15]
    return outs[0] if outs else r["location"]


def clock_in_store(target):
    """each current shift's store, keyed by colleague, date and start"""
    rows, pun = wfm_shifts(target), punches_by_employee(target)
    return {(r["employee_no"], r["shift_date"], r["shift_start"]): worked_store(r, pun) for r in current_shifts(rows, pun)}


def ask_b(target, mode="golden"):
    """paid hours a week worked at each store over the 25 weeks: each shift worked once, at the hours it was paid,
    at the store its colleague clocked in at"""
    rows = wfm_shifts(target)
    pun = punches_by_employee(target)
    tot = {}
    if mode == "punch_hours":
        for e, ps in pun.items():
            ps = sorted(ps)
            for (t0, k0, s0), (t1, k1, s1) in zip(ps, ps[1:]):
                if k0 == "IN" and k1 == "OUT":
                    tot[s0] = tot.get(s0, 0.0) + (t1 - t0).total_seconds() / 3600.0
        return {s: v / 25.0 for s, v in tot.items()}
    if mode == "golden":
        use = current_shifts(rows, pun)
    elif mode == "by_file":
        use = current_shifts(rows, pun, by="file")
    elif mode in ("dedup_last", "dedup_first", "dedup_max"):
        g = {}
        for r in rows:
            g.setdefault((r["employee_no"], r["shift_date"]), []).append(r)
        if mode == "dedup_last":
            use = [max(v, key=lambda r: r["i"]) for v in g.values()]
        elif mode == "dedup_first":
            use = [min(v, key=lambda r: r["i"]) for v in g.values()]
        else:
            use = [max(v, key=lambda r: (r["paid"], r["i"])) for v in g.values()]
    else:
        use = rows
    for r in use:
        s = r["location"] if mode == "location" else worked_store(r, pun)
        tot[s] = tot.get(s, 0.0) + r["paid"]
    return {s: v / 25.0 for s, v in tot.items()}


# ----------------------------------------------------------------- ask C
def _mins(t):
    h, m, s = (int(x) for x in t.split(":"))
    return 60 * h + m


def ask_c(target, mode="golden", day=P.BUS_DAY):
    """bus departures on the day, by the clock, from stops within 400 metres of each site: the day's own services up
    to midnight and the previous day's services after it (a GTFS service day runs past 24:00), no departure from a
    stop where the bus only sets down, and each bus once: both operators' feeds carry the jointly run routes'
    journeys, so a journey both list is one bus, while a second bus one feed lists on the same timetable (a school-day
    duplicate) is another"""
    import collections
    import geo
    sites = sheet(os.path.join(target, P.FILES["shortlist"]), "Shortlist")
    pc = {r["pcds"]: (float(r["oseast1m"]), float(r["osnrth1m"])) for r in rd(os.path.join(target, P.FILES["onspd"]))}
    site_ll = {r["site"]: geo.bng_to_wgs84(*pc[r["postcode"]]) for r in sites}
    per_feed = []
    late = mode in ("golden", "nodedup", "all_calls", "set")
    days = [(day, "today")] + ([(day - dt.timedelta(days=1), "eve")] if late else [])
    for key in ("gtfs_lb", "gtfs_pdc"):
        cnt = collections.Counter()
        with zipfile.ZipFile(os.path.join(target, P.FILES[key])) as z:
            R = lambda n: list(csv.DictReader(io.TextIOWrapper(z.open(n), encoding="utf-8")))
            cal = R("calendar.txt")
            cds = R("calendar_dates.txt")
            routes = {r["route_id"]: r["route_short_name"] for r in R("routes.txt")}
            trips = R("trips.txt")
            stops = {r["stop_id"]: (float(r["stop_lat"]), float(r["stop_lon"])) for r in R("stops.txt")}
            near = {k: {sid for sid, ll in stops.items() if geo.haversine_m(ll[0], ll[1], *site_ll[k]) <= P.BUS_RADIUS}
                    for k in site_ll}
            st = R("stop_times.txt")
            for d, which in days:
                ymd, dow = d.strftime("%Y%m%d"), d.strftime("%A").lower()
                on = {r["service_id"] for r in cal if r[dow] == "1" and r["start_date"] <= ymd <= r["end_date"]}
                for r in cds:
                    if r["date"] == ymd:
                        (on.add if r["exception_type"] == "1" else on.discard)(r["service_id"])
                tr = {r["trip_id"]: (routes[r["route_id"]], r["direction_id"]) for r in trips if r["service_id"] in on}
                for r in st:
                    if r["trip_id"] not in tr or (r["pickup_type"] == "1" and mode != "all_calls"):
                        continue
                    m = _mins(r["departure_time"])
                    if which == "today" and mode != "naive" and m >= 1440:
                        continue
                    if which == "eve" and m < 1440:
                        continue
                    for k in site_ll:
                        if r["stop_id"] in near[k]:
                            cnt[(k, tr[r["trip_id"]][0], tr[r["trip_id"]][1], r["stop_id"], m, which)] += 1
        per_feed.append(cnt)
    out = {k: 0 for k in site_ll}
    for key in set(per_feed[0]) | set(per_feed[1]):
        if mode == "nodedup":
            n = per_feed[0][key] + per_feed[1][key]
        elif mode in ("set", "naive", "half"):
            n = 1
        else:
            n = max(per_feed[0][key], per_feed[1][key])
        out[key[0]] += n
    return out
