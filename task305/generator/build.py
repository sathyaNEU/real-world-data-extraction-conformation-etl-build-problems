"""Cut the task305 pack from the world, the records and the operational records."""
import csv
import datetime as dt
import io
import math
import os
import zipfile

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import docs
import geo
import ops as O
import params as P
import postprocess
import records as R
import world as Wd

FILES = P.FILES
PACK = list(FILES)
DISTRACTORS = ("aadf", "survey")
CONTEXT = ("cards",)


def _csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)


def _xlsx(path, sheets, made, saved, author):
    wb = Workbook()
    for i, (title, header, rows, widths, fmts) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DCE6DF")
            c.alignment = Alignment(vertical="top", wrap_text=True)
        for r in rows:
            ws.append(list(r))
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        for col, f in (fmts or {}).items():
            j = header.index(col) + 1
            for k in range(2, len(rows) + 2):
                ws.cell(row=k, column=j).number_format = f
        ws.freeze_panes = "A2"
    wb.properties.creator = author
    wb.properties.lastModifiedBy = author
    wb.save(path)
    postprocess.scrub(path, P.ORG, saved[:10])
    postprocess.finish_office(path, made, saved, P.ORG)


def hhmm(t):
    return t.strftime("%Y-%m-%d %H:%M")


def gtfs_time(m):
    m = int(m)
    return "%02d:%02d:00" % (m // 60, m % 60)


def locality(G):
    """the town a point lies in, or else the parish of the nearest village seat; a site's own surroundings take the
    town its address gives"""
    def at(x, y):
        for st in P.SITES:
            if math.hypot(st[3] - x, st[4] - y) <= 1.2:
                return st[2]
        t = min(P.TOWNS, key=lambda t: (t[1] - x) ** 2 + (t[2] - y) ** 2)
        if math.hypot(t[1] - x, t[2] - y) <= 2.5:
            return t[0]
        return min(G["seats"], key=lambda q: (q[1] - x) ** 2 + (q[2] - y) ** 2)[0]
    return at


TOWN_STOPS = ["Market Place", "Bus Station", "High Street", "Station Road", "Church", "Health Centre", "Library",
              "Park", "Fire Station", "School", "Leisure Centre", "Cemetery", "Post Office", "Recreation Ground"]
VILLAGE_STOPS = ["Church", "The Green", "Post Office", "Crossroads", "Village Hall", "Chapel", "Farm Lane", "Bridge",
                 "Mill", "School", "Allotments", "Main Street"]


def survey_rows(own, ds):
    """two waves a store: each store's scores move a little between waves around its own level"""
    out = []
    for s in own:
        sat, nps = ds.uniform(7.3, 8.4), ds.uniform(14, 36)
        avail, queue, staff, value = ds.uniform(80, 92), ds.uniform(72, 90), ds.uniform(83, 94), ds.uniform(63, 82)
        for w in ("2025 spring", "2025 autumn"):
            out.append([s["id"], w, int(ds.integers(180, 420)), "%.1f" % (sat + ds.normal(0, .12)),
                        int(round(nps + ds.normal(0, 3))), int(round(avail + ds.normal(0, 1.5))),
                        int(round(queue + ds.normal(0, 2))), int(round(staff + ds.normal(0, 1.5))),
                        int(round(value + ds.normal(0, 2)))])
    return out


def write_gtfs(path_lb, path_pdc, feeds, stops, at):
    rs = np.random.default_rng(P.OPS_SEED + 9)
    # both operators publish on the region's NaPTAN stops: the same ATCO code, SMS code and name for a stop
    lb_stop = {}
    towns = {t[0] for t in P.TOWNS}
    count = {}
    for i, k in enumerate(sorted(stops)):
        x, y, _ = stops[k]
        lat, lon = geo.bng_to_wgs84(P.E0 + 1000 * x, P.N0 + 1000 * y)
        place = at(x, y)
        pool = TOWN_STOPS if place in towns else VILLAGE_STOPS
        n = count.get(place, 0)
        count[place] = n + 1
        name = pool[n % len(pool)] if n < len(pool) else "%s %s" % (pool[n % len(pool)], "North" if n // len(pool) == 1 else "South")
        ind = ["opp", "adj", "o/s", "nr", "Stop A", "Stop B", "N-bound", "S-bound", "E-bound", "W-bound"][int(rs.integers(0, 10))]
        name = "%s, %s (%s)" % (place, name, ind) if name != "Bus Station" else "%s, Bus Station" % place
        code = "lq" + "".join(rs.choice(list("abcdefghjkmnprstuvwxyz"), 6))
        atco = "2709" + "%08d" % int(rs.integers(10000000, 99999999))
        while any(v[0] == atco for v in lb_stop.values()):
            atco = "2709" + "%08d" % int(rs.integers(10000000, 99999999))
        lb_stop[k] = (atco, code, name, round(lat, 6), round(lon, 6))
    pdc_stop = dict(lb_stop)
    for owner, path, smap, noc, agency in (("LB", path_lb, lb_stop, "SLMS", "Sallowmere Motor Services"),
                                            ("PDC", path_pdc, pdc_stop, "WDCO", "Wrentoft & District Coaches")):
        trips = feeds[owner]["trips"]
        used = sorted({k for t in trips for k in t["keys"]})
        files = {}
        files["agency.txt"] = [["agency_id", "agency_name", "agency_url", "agency_timezone", "agency_lang", "agency_phone",
                                "agency_noc"], [noc, agency, "https://www.%s.co.uk" % agency.lower().replace(" & ", "and").replace(" ", ""),
                                                "Europe/London", "EN", "01632 960%03d" % int(rs.integers(0, 1000)), noc]]
        rmap = {}
        rows = [["route_id", "agency_id", "route_short_name", "route_long_name", "route_type"]]
        for t in trips:
            if t["num"] not in rmap:
                rmap[t["num"]] = "%s%d" % (noc[:2], 400 + len(rmap) * 7)
                ends = {"4": "Quellby - Darrowfield - Sallowmere", "41": "Garretstead - Brackenmere - Darrowfield - Bellingtoft - Wrentoft - Skerryby",
                        "7": "Darrowfield Town Circular", "12": "Brackenmere - Corrinby", "15": "Darrowfield - Raddleford - Loderwell",
                        "22": "Bellingtoft - Sallowmere - Sedgeton", "9": "Wrentoft - Withybeck", "S41": "Wrentoft - Withybeck School"}
                rows.append([rmap[t["num"]], noc, t["num"], ends[t["num"]], 3])
        files["routes.txt"] = rows
        # each operator's own registrations: Sallowmere Motor Services' run from the October timetable change to the
        # end of March, Wrentoft & District Coaches' from their 12 October change to the start of April
        if owner == "LB":
            sid = {"wk": "2201", "sat": "2202", "sun": "2203", "mtl": "2207", "fl": "2208", "sl": "2209", "sch": "2230"}
            v0, v1 = "20261005", "20270328"
        else:
            sid = {"wk": "MF-W26", "sat": "SA-W26", "sun": "SU-W26", "mtl": "MTL-W26", "fl": "FL-W26", "sl": "SL-W26",
                   "sch": "SCH-2627"}
            v0, v1 = "20261012", "20270404"
        cal = [["service_id", "monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday", "start_date",
                "end_date"],
               [sid["wk"], 1, 1, 1, 1, 1, 0, 0, v0, v1], [sid["sat"], 0, 0, 0, 0, 0, 1, 0, v0, v1],
               [sid["sun"], 0, 0, 0, 0, 0, 0, 1, v0, v1]]
        late = any(t["day"] in ("mtl", "fl", "sl") for t in trips)
        if late:
            cal += [[sid["mtl"], 1, 1, 1, 1, 0, 0, 0, v0, v1], [sid["fl"], 0, 0, 0, 0, 1, 0, 0, v0, v1],
                    [sid["sl"], 0, 0, 0, 0, 0, 1, 0, v0, v1]]
        school = any(t["school"] for t in trips)
        if school and owner == "PDC":
            # Wrentoft & District Coaches register their school journeys for the school year, holidays taken out
            cal.append([sid["sch"], 1, 1, 1, 1, 1, 0, 0, "20260907", "20270723"])
        files["calendar.txt"] = cal
        cd = [["service_id", "date", "exception_type"]]
        if owner == "LB":
            for d in ("20261225", "20261228", "20270101"):
                cd.append([sid["wk"], d, 2]); cd.append([sid["sun"], d, 1])
            cd.append([sid["sat"], "20261226", 2])
            if late:
                for d in ("20261224", "20261231"):
                    cd.append([sid["mtl"], d, 2])
                for d in ("20261225", "20270101"):
                    cd.append([sid["fl"], d, 2])
                cd.append([sid["sl"], "20261226", 2])
        else:
            for d in ("20261225", "20261226", "20261228", "20270101"):
                cd.append([sid["wk"] if d != "20261226" else sid["sat"], d, 2])
            cd.append([sid["sun"], "20261228", 1]); cd.append([sid["sun"], "20270101", 1])
            if late:
                for d in ("20261224", "20261231"):
                    cd.append([sid["mtl"], d, 2])
                cd.append([sid["fl"], "20261225", 2]); cd.append([sid["fl"], "20270101", 2])
                cd.append([sid["sl"], "20261226", 2])
        holidays = [(dt.date(2026, 10, 26), dt.date(2026, 10, 30)), (dt.date(2026, 12, 21), dt.date(2027, 1, 4)),
                    (dt.date(2027, 2, 15), dt.date(2027, 2, 19)), (dt.date(2027, 3, 29), dt.date(2027, 4, 9)),
                    (dt.date(2027, 5, 31), dt.date(2027, 6, 4))]
        if school and owner == "LB":
            # Sallowmere Motor Services list each school day of the timetable's run
            d = dt.date(2026, 10, 5)
            while d <= dt.date(2027, 3, 26):
                if d.weekday() < 5 and not any(h0 <= d <= h1 for h0, h1 in holidays):
                    cd.append([sid["sch"], d.strftime("%Y%m%d"), 1])
                d += dt.timedelta(days=1)
        elif school:
            for h0, h1 in holidays:
                d = h0
                while d <= h1:
                    if d.weekday() < 5:
                        cd.append([sid["sch"], d.strftime("%Y%m%d"), 2])
                    d += dt.timedelta(days=1)
        files["calendar_dates.txt"] = cd
        tr = [["route_id", "service_id", "trip_id", "trip_headsign", "direction_id", "block_id", "wheelchair_accessible",
               "vehicle_journey_code"]]
        st = [["trip_id", "arrival_time", "departure_time", "stop_id", "stop_sequence", "pickup_type", "drop_off_type",
               "timepoint"]]
        order = sorted(trips, key=lambda t: (t["num"], t["direction"], t["day"], t["t0"]))
        for i, t in enumerate(order):
            trip_id = "VJ_%s_%06d" % (noc, 104000 + 17 * i + int(rs.integers(0, 9)))
            service = sid["sch"] if t["school"] else sid[t["day"]]
            head = at(*stops[t["keys"][-1]][:2])
            tr.append([rmap[t["num"]], service, trip_id, head, t["direction"], t["block"],
                       1, "%d" % (2000 + 2 * i + (1 if owner == "PDC" else 0))])
            for j, (k, m) in enumerate(zip(t["keys"], t["times"])):
                last = j == len(t["keys"]) - 1
                st.append([trip_id, gtfs_time(m), gtfs_time(m), smap[k][0], j + 1, 1 if last else 0, 1 if j == 0 else 0,
                           1 if (j % 4 == 0 or last) else 0])
        files["trips.txt"] = tr
        files["stop_times.txt"] = st
        files["stops.txt"] = [["stop_id", "stop_code", "stop_name", "stop_lat", "stop_lon", "wheelchair_boarding",
                               "location_type", "parent_station"]] + \
                              [[smap[k][0], smap[k][1], smap[k][2], "%.6f" % smap[k][3], "%.6f" % smap[k][4], 0, 0, ""]
                               for k in used]
        files["feed_info.txt"] = [["feed_publisher_name", "feed_publisher_url", "feed_lang", "feed_start_date",
                                   "feed_end_date", "feed_version"],
                                  ["Bus Open Data Service", "https://data.bus-data.dft.gov.uk", "EN", v0, v1,
                                   "%s-20261002" % noc]]
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            for name in ("agency.txt", "calendar.txt", "calendar_dates.txt", "feed_info.txt", "routes.txt", "stop_times.txt",
                         "stops.txt", "trips.txt"):
                buf = io.StringIO()
                w = csv.writer(buf, lineterminator="\r\n")
                w.writerows(files[name])
                zi = zipfile.ZipInfo(name, date_time=(2026, 10, 2, 3, 15, 0))
                zi.compress_type = zipfile.ZIP_DEFLATED
                z.writestr(zi, buf.getvalue().encode("utf-8"))
    return lb_stop, pdc_stop


def build(target):
    os.makedirs(target, exist_ok=True)
    F = {k: os.path.join(target, v) for k, v in FILES.items()}
    T = Wd.truth()
    G = R.geography()
    oas = G["oas"]
    # census, centroids, spend
    _csv(F["households"], ["date", "geography", "geography code", "Number of households"],
         [[2021, o["oa"], o["oa"], o["hh"]] for o in sorted(oas, key=lambda o: o["oa"])])
    gs = np.random.default_rng(P.SEED + 41)
    _csv(F["pwc"], ["FID", "OA21CD", "x", "y"],
         [[i + 1, o["oa"], "%.3f" % (P.E0 + 1000 * o["x"]), "%.3f" % (P.N0 + 1000 * o["y"])]
          for i, o in enumerate(sorted(oas, key=lambda o: o["oa"]))])
    _csv(F["spend"], ["OA21CD", "AvgWeeklyGrocerySpendPerHousehold_GBP_2026"],
         [[o["oa"], "%.2f" % o["spend"]] for o in sorted(oas, key=lambda o: o["oa"])])
    _csv(F["onspd"], ["pcd", "pcd2", "pcds", "dointr", "doterm", "usertype", "oseast1m", "osnrth1m", "osgrdind", "oa21",
                      "lsoa21", "oslaua"], R.onspd_rows(G))
    rl, smv = R.rating_list(G)
    with open(F["voa_list"], "w", encoding="utf-8", newline="") as fh:
        for r in rl:
            fh.write("*".join("" if v is None else str(v) for v in r) + "\r\n")
    _csv(F["voa_smv"], ["uarn", "assessment_reference", "scat_code", "main_measure", "unit_of_measurement"], smv)
    # the card feed; the estate's sales are what the feed's card spend implies at each store's card share
    cf, card_tot = R.card_feed(G, T)
    sales = {sid: card_tot[sid] / T["card"][sid] for sid in card_tot}
    # the estate
    own = [s for s in G["stores"] if s["kind"] == "own"]
    opened = {o[0]: o[5] for o in P.OWN}
    _xlsx(F["estate"], [("Stores", ["store_no", "store", "address", "postcode", "opened", "floor_area_m2",
                                    "sales_fy2025_26_gbp", "card_share_of_sales"],
                         [[s["id"], "Marrable %s" % s["town"], "%s, %s" % (R.STREETS[int(s["id"]) % len(R.STREETS)], s["town"]),
                           s["pc"], opened[s["id"]], s["area"], int(round(sales[s["id"]])), T["card"][s["id"]]]
                          for s in own], [10, 22, 34, 12, 9, 15, 19, 18],
                         {"sales_fy2025_26_gbp": "#,##0", "card_share_of_sales": "0.000"})],
          "2021-05-04T09:12:00Z", "2026-10-07T16:44:00Z", P.PEOPLE["property"])
    _csv(F["cards"], ["store_no", "month", "cardholder_home_sector", "cards", "transactions", "spend_gbp"],
         [[r[0], r[1], r[2], r[3], r[4], "%.2f" % r[5]] for r in cf])
    # the shortlist
    plan = {"S1": "full permission granted 14 May 2025", "S2": "full permission granted 2 July 2024",
            "S3": "full permission granted 19 March 2025", "S4": "outline permission; reserved matters submitted",
            "S5": "full permission granted 30 October 2024", "S6": "full permission granted 11 February 2025"}
    parking = {"S1": 182, "S2": 165, "S3": 140, "S4": 151, "S5": 158, "S6": 143}
    _xlsx(F["shortlist"], [("Shortlist", ["site", "address", "town", "postcode", "proposed_floor_area_m2", "rent_gbp",
                                          "business_rates_gbp", "staffing_gbp", "other_running_gbp", "running_costs_gbp",
                                          "parking_spaces", "planning"],
                            [[s[0], s[1], s[2], G["sites"][s[0]], s[5]] + list(s[6]) + [sum(s[6]), parking[s[0]], plan[s[0]]]
                             for s in P.SITES], [6, 16, 14, 11, 12, 11, 12, 12, 12, 13, 9, 40],
                            {k: "#,##0" for k in ("rent_gbp", "business_rates_gbp", "staffing_gbp", "other_running_gbp",
                                                  "running_costs_gbp")})],
          "2026-06-22T10:05:00Z", "2026-10-06T11:20:00Z", P.PEOPLE["property"])
    docs.model_sheet(F["model"])
    postprocess.scrub(F["model"], "Microsoft Word 2019", "2025-06-27T14:40:00")
    docs.notes(F["notes"])
    # ask A
    rows, recs, at = O.tms(T, G)
    _csv(F["tms"], ["trip_id", "trailer_id", "depart_dc_utc", "stop_seq", "store_no", "arrive_utc", "depart_utc", "stop_type",
                    "status", "cases", "cages"],
         [[r[0], r[1], hhmm(r[2]), r[3], r[4], hhmm(r[5]), hhmm(r[6]), r[7], r[8], r[9], r[10]] for r in rows])
    _csv(F["receipts"], ["receipt_no", "store_no", "received_local", "trailer_id", "delivery_note_no", "cases"],
         [[r[0], r[1], hhmm(r[2]), r[3], r[4], r[5]] for r in recs])
    # ask B
    shifts, punches, register, st = O.staffing(T)
    _csv(F["wfm"], ["employee_no", "location", "date", "role", "paid_hours"],
         [[r[0], r[1], r[2].isoformat(), r[3], "%.2f" % r[4]] for r in shifts])
    _csv(F["punches"], ["employee_no", "terminal_id", "punch_time", "punch"],
         [[p[0], p[1], hhmm(p[2]), p[3]] for p in punches])
    _csv(F["terminals"], ["terminal_id", "store_no", "door", "installed"], register)
    # ask C
    feeds, stops, bt, near, dups, seam = O.buses(T)
    write_gtfs(F["gtfs_lb"], F["gtfs_pdc"], feeds, stops, locality(G))
    # distractors: the DfT's count points, the customer survey
    ds = np.random.default_rng(P.SEED + 53)
    roads = ["A153", "A158", "A1121", "A16", "A155", "B1191", "B1183", "A1084", "B1192", "B1195", "A631", "B1225"]
    aadf = []
    for i in range(64):
        x, y = ds.uniform(2, 70), ds.uniform(2, 44)
        lat, lon = geo.bng_to_wgs84(P.E0 + 1000 * x, P.N0 + 1000 * y)
        major = i % 3 != 0
        cars = int(ds.uniform(6000, 17000) if major else ds.uniform(900, 4200))
        method = "Counted" if ds.uniform() < .45 else "Estimated"
        aadf.append([946000 + 113 * i, 2025, "East Midlands", "Sallowdale", roads[i % len(roads)],
                     "PA" if major else "MB", int(P.E0 + 1000 * x), int(P.N0 + 1000 * y), "%.5f" % lat, "%.5f" % lon,
                     "%.1f" % ds.uniform(0.8, 9.5), method,
                     ("Manual count" if ds.uniform() < .7 else "Automatic counter") if method == "Counted" else
                     ("Estimated using previous year's AADF on this link" if ds.uniform() < .75 else
                      "Estimated from nearby links"),
                     int(cars * .006), int(cars * .009), cars, int(cars * .007), int(cars * .19), int(cars * .045),
                     int(cars * 1.25)])
    _csv(F["aadf"], ["count_point_id", "year", "region_name", "local_authority_name", "road_name", "road_category", "easting",
                     "northing", "latitude", "longitude", "link_length_km", "estimation_method", "estimation_method_detailed",
                     "pedal_cycles", "two_wheeled_motor_vehicles", "cars_and_taxis", "buses_and_coaches", "lgvs", "all_hgvs",
                     "all_motor_vehicles"], aadf)
    _csv(F["survey"], ["store_no", "wave", "responses", "overall_satisfaction_0_10", "net_promoter_score", "availability_pct",
                       "queue_time_pct", "staff_helpfulness_pct", "value_for_money_pct"],
         survey_rows(own, ds))
    return dict(T=T, G=G, rl=rl, smv=smv, cf=cf, shifts=shifts, punches=punches, register=register, st=st, tms=rows,
                recs=recs, at=at, feeds=feeds, stops=stops, bt=bt, near=near, dups=dups, seam=seam)


if __name__ == "__main__":
    import sys
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "target")
    build(out)
    print("built", out)
