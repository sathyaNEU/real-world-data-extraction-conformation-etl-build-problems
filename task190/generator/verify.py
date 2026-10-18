"""Independent verifier. It reads ONLY the shipped bundle under target/ and the goldens under
golden/, imports nothing from the generator, and recomputes every figure the deliverables state,
including the plan targets, which it reads out of the plan PDF."""
import collections
import csv
import math
import os
import re
import sys

TASK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TGT = os.path.join(TASK, "target")
GLD = os.path.join(TASK, "golden")
BAR = 1.959963985
FAILS = []
N = [0]


def ok(label, cond, detail=""):
    N[0] += 1
    if not cond:
        FAILS.append(label)
        print("  FAIL  %-54s %s" % (label, detail))
    else:
        print("  ok    %-54s %s" % (label, detail))


def read(name):
    with open(os.path.join(TGT, name), newline="", encoding="utf-8") as fh:
        rd = csv.reader(fh)
        hdr = next(rd)
        return hdr, [r for r in rd]


def pdf_text(name):
    from pypdf import PdfReader
    return "\n".join(p.extract_text() or "" for p in PdfReader(os.path.join(TGT, name)).pages)


def main():
    # ---- the survey side: strata recovered from the micro-data alone
    hdr, md = read("bws_2026_microdata.csv")
    cols = [c for c in hdr if c.startswith("arisings_t_")]
    keys = [c[len("arisings_t_"):] for c in cols]
    ci = {c: hdr.index(c) for c in hdr}
    by_h = collections.defaultdict(list)
    for r in md:
        by_h[r[ci["stratum_code"]]].append(r)
    strata = {}
    for h, rs in by_h.items():
        w = float(rs[0][ci["design_weight"]])
        assert all(abs(float(x[ci["design_weight"]]) - w) < 1e-9 for x in rs)
        r_h = len(rs)
        strata[h] = {"w": w, "r": r_h, "N": int(round(w * r_h)), "rows": rs}
    ok("the micro-data carries one design weight per stratum", len(strata) == 4,
       "strata %s, frame %d, usable returns %d"
       % (sorted(strata), sum(s["N"] for s in strata.values()), len(md)))

    est, var = {}, {}
    for k, c in zip(keys, cols):
        j = ci[c]
        tot, V = 0.0, 0.0
        for h, s in strata.items():
            vals = [float(x[j]) for x in s["rows"]]
            tot += s["w"] * sum(vals)
            if s["r"] >= s["N"]:
                continue
            m = sum(vals) / len(vals)
            s2 = sum((v - m) ** 2 for v in vals) / (len(vals) - 1)
            V += s["N"] ** 2 * (1.0 - s["r"] / s["N"]) * s2 / s["r"]
        est[k], var[k] = tot, V

    # ---- the register side
    kinds, revision = {}, {}
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(TGT, "treatment_operation_codes.xlsx"), data_only=True)
    ws = wb["Operation codes"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            kinds[row[0]] = row[1]
    _h, cwr = read("waste_code_crosswalk.csv")
    code2stream = {r[0]: r[1] for r in cwr}
    _h, facr = read("facility_register.csv")
    fac = {r[0] for r in facr}

    def legs(name):
        hdr, rs = read(name)
        return rs

    ch = collections.defaultdict(list)
    for nm in ("waste_movement_legs_2026_h1.csv", "waste_movement_legs_2026_h2.csv"):
        for r in legs(nm):
            ch[r[0]].append(r)
    for c in ch:
        ch[c].sort(key=lambda r: int(r[1]))
    child_of = {}
    for c, rs in ch.items():
        if rs[0][10]:
            child_of.setdefault(rs[0][10], c)
    po = {c for c, rs in ch.items()
          if [x for x in rs if x[2] == "PRODUCER_DESPATCH"][0][4] not in fac}

    def terminal(c):
        conf = [x for x in ch[c] if x[2] == "TREATMENT_CONFIRMED"]
        op = conf[0][9] if conf else ""
        if kinds.get(op) == "interim":
            k = child_of.get(c)
            cc = [x for x in ch[k] if x[2] == "TREATMENT_CONFIRMED"] if k else []
            op = cc[0][9] if cc else ""
        return op

    rec = collections.Counter()
    desp = collections.Counter()
    conf_t = collections.Counter()
    for c in po:
        d = [x for x in ch[c] if x[2] == "PRODUCER_DESPATCH"][0]
        if not d[3].startswith("2026"):
            continue
        s, t = code2stream[d[7]], float(d[8])
        desp[s] += t
        op = terminal(c)
        if kinds.get(op) in ("recovery", "disposal"):
            conf_t[s] += t
            if kinds[op] == "recovery":
                rec[s] += t

    for c in (rec, desp, conf_t):
        for k, v in c.items():
            assert abs(v - round(v)) < 1e-6, (k, v)
    rec = {k: int(round(v)) for k, v in rec.items()}
    desp = {k: int(round(v)) for k, v in desp.items()}
    conf_t = {k: int(round(v)) for k, v in conf_t.items()}
    ok("every stream's register total is a whole tonne once the weighbridge noise is summed",
       True, "net tonnes are recorded to 20 kg and sum to whole tonnes by stream and stage")

    # ---- the plan targets, read out of the plan PDF
    txt = pdf_text("national_waste_management_plan_targets.pdf")
    tgt = {}
    names = {}
    lines = [l.strip() for l in txt.splitlines()]
    for i, l in enumerate(lines):
        if l.endswith("wastes") and i + 1 < len(lines):
            m = re.match(r"^([0-9]{1,2}\.[0-9]) per cent$", lines[i + 1])
            if m:
                names[l] = float(m.group(1))
    ok("the plan PDF yields a 2026 target for each of the eight streams", len(names) == 8,
       "%d targets read from the PDF" % len(names))

    # match the plan's stream names to the crosswalk's statistical_stream values
    name2key = {r[2]: r[1] for r in cwr}
    for nm, v in names.items():
        if nm in name2key:
            tgt[name2key[nm]] = v / 100.0
    ok("every plan target maps to a statistical stream key", len(tgt) == 8, "%d mapped" % len(tgt))

    # ---- the call
    res = {}
    for k in keys:
        X = est[k]
        rate = rec[k] / X
        se = rate * math.sqrt(var[k]) / X
        res[k] = {"rec": rec[k], "X": X, "rate": rate, "se": se, "tgt": tgt[k],
                  "short": tgt[k] - rate, "z": (tgt[k] - rate) / se,
                  "lo": rate - BAR * se, "hi": rate + BAR * se}
    cl = [k for k in keys if res[k]["z"] > BAR]
    ok("exactly one stream is established below target", len(cl) == 1, "established: %s" % cl)
    ok("the established stream is metal and metal-bearing wastes", cl == ["metal"],
       "z %.2f, rate %.1f per cent, target %.1f per cent, shortfall %.1f pp"
       % (res["metal"]["z"], res["metal"]["rate"] * 100, res["metal"]["tgt"] * 100,
          res["metal"]["short"] * 100))
    loud = max(keys, key=lambda k: res[k]["short"])
    ok("the largest raw shortfall is a stream that is not established",
       loud == "mixed_ci" and res[loud]["z"] < BAR,
       "%s, shortfall %.1f pp, z %.2f" % (loud, res[loud]["short"] * 100, res[loud]["z"]))

    # ---- the goldens, cell by cell
    wb = load_workbook(os.path.join(GLD, "recovery_attainment_2026.xlsx"), data_only=True)
    ws = wb["Attainment"]
    seen = 0
    for row in ws.iter_rows(min_row=6, max_row=13, values_only=True):
        if not row[0]:
            continue
        k = name2key[row[0]]
        r = res[k]
        bad = []
        for label, got, want in (("recovered", row[1], r["rec"]),
                                 ("arisings", row[2], int(round(r["X"]))),
                                 ("rate", row[3], round(r["rate"] * 100, 1)),
                                 ("se", row[4], round(r["se"] * 100, 2)),
                                 ("lo", row[5], round(r["lo"] * 100, 1)),
                                 ("hi", row[6], round(r["hi"] * 100, 1)),
                                 ("target", row[7], round(r["tgt"] * 100, 1)),
                                 ("shortfall", row[8], round(r["short"] * 100, 1)),
                                 ("z", row[9], round(r["z"], 2)),
                                 ("verdict", row[10], "yes" if r["z"] > BAR else "no")):
            if got != want:
                bad.append("%s %r != %r" % (label, got, want))
        ok("Attainment row recomputes: %s" % k, not bad, "; ".join(bad))
        seen += 1
    ok("the Attainment sheet carries all eight streams", seen == 8, "%d rows" % seen)

    ws = wb["Stage chain"]
    for row in ws.iter_rows(min_row=4, max_row=11, values_only=True):
        if not row[0]:
            continue
        k = name2key[row[0]]
        bad = []
        for label, got, want in (("arisings", row[1], int(round(est[k]))),
                                 ("despatched", row[2], desp[k]),
                                 ("confirmed", row[3], conf_t[k]),
                                 ("recovered", row[4], rec[k]),
                                 ("p12", row[5], round(100.0 * desp[k] / est[k], 1)),
                                 ("p23", row[6], round(100.0 * conf_t[k] / desp[k], 1)),
                                 ("p34", row[7], round(100.0 * rec[k] / conf_t[k], 1))):
            if got != want:
                bad.append("%s %r != %r" % (label, got, want))
        ok("Stage chain row recomputes: %s" % k, not bad, "; ".join(bad))

    # ---- the asks
    _h, cwe = read("enterprise_permit_crosswalk.csv")
    region_of = {r[0]: r[2] for r in cwe}
    _h, lg = read("bws_2026_response_log.csv")
    days, ret, seenv = collections.Counter(), collections.Counter(), set()
    for r in lg:
        reg = region_of[r[1]]
        fd = float(r[7]) / (8.0 if r[8] == "h" else 1.0)
        if r[5]:
            if r[5] in seenv:
                continue
            seenv.add(r[5])
        days[reg] += fd
    done = set()
    for r in lg:
        if r[9] == "usable return" and r[0] not in done:
            done.add(r[0])
            ret[region_of[r[1]]] += 1
    ok("the usable returns counted by case tie to the micro-data's records",
       sum(ret.values()) == len(md), "%d usable returns over five regions, %d micro-data records"
       % (sum(ret.values()), len(md)))
    ws = wb["Collection effort"]
    for row in ws.iter_rows(min_row=5, max_row=9, values_only=True):
        if not row[0]:
            continue
        reg = row[0]
        bad = []
        for label, got, want in (("days", row[1], round(days[reg], 2)),
                                 ("returns", row[2], ret[reg]),
                                 ("per day", row[3], round(ret[reg] / days[reg], 2))):
            if got != want:
                bad.append("%s %r != %r" % (label, got, want))
        ok("Collection effort row recomputes: %s" % reg, not bad, "; ".join(bad))

    ch2 = collections.defaultdict(list)
    for r in legs("waste_movement_legs_2025.csv"):
        ch2[r[0]].append(r)
    for c in ch2:
        ch2[c].sort(key=lambda r: int(r[1]))
    co2 = {}
    for c, rs in ch2.items():
        if rs[0][10]:
            co2.setdefault(rs[0][10], c)
    po2 = {c for c, rs in ch2.items()
           if [x for x in rs if x[2] == "PRODUCER_DESPATCH"][0][4] not in fac}
    a2 = collections.Counter()
    for c in po2:
        d = [x for x in ch2[c] if x[2] == "PRODUCER_DESPATCH"][0]
        if not d[3].startswith("2025") or (d[10] and d[10] in po2):
            continue
        conf = [x for x in ch2[c] if x[2] == "TREATMENT_CONFIRMED"]
        op = conf[0][9] if conf else ""
        if kinds.get(op) == "interim":
            k2 = co2.get(c)
            cc = [x for x in ch2[k2] if x[2] == "TREATMENT_CONFIRMED"] if k2 else []
            op = cc[0][9] if cc else ""
        if kinds.get(op) == "recovery":
            a2[code2stream[d[7]]] += float(d[8])
    a2 = {k: int(round(v)) for k, v in a2.items()}
    ws = wb["Recovery 2025"]
    for row in ws.iter_rows(min_row=4, max_row=11, values_only=True):
        if not row[0]:
            continue
        k = name2key[row[0]]
        ok("Recovery 2025 recomputes: %s" % k, row[1] == a2[k], "%s t" % a2[k])

    _h, mret = read("municipal_waste_returns_2026.csv")
    _h, mrec = read("municipal_facility_receipts_2026.csv")
    muni = {r[0] for r in mret}
    members = collections.defaultdict(set)
    col, rcv = collections.Counter(), collections.Counter()
    tot_by = collections.defaultdict(float)
    for r in mret:
        v = float(r[3]) / (10.0 if r[4] == "100kg" else 1.0)
        tot_by[(r[0], r[1])] += v
    # an entity whose monthly total equals the sum of three other entities' totals in its region,
    # month by month, is reporting on their behalf
    bym = collections.defaultdict(dict)
    for r in mret:
        v = float(r[3]) / (10.0 if r[4] == "100kg" else 1.0)
        bym[(r[0], r[1])][r[2]] = v
    auth = set()
    for (ent, reg), series in bym.items():
        others = [(e2, s2) for (e2, r2), s2 in bym.items() if r2 == reg and e2 != ent]
        for i in range(len(others)):
            for j in range(i + 1, len(others)):
                for l in range(j + 1, len(others)):
                    trio = [others[i][1], others[j][1], others[l][1]]
                    if all(abs(series[m] - sum(t[m] for t in trio)) < 0.05 for m in series):
                        auth.add(ent)
    ok("the verifier finds the reporting-on-behalf entities from the data alone", len(auth) == 5,
       "%d found: %s" % (len(auth), sorted(auth)[:2]))
    for (ent, reg), series in bym.items():
        if ent in auth:
            continue
        col[reg] += sum(series.values())
    for r in mrec:
        rcv[r[3]] += float(r[4])
    ws = wb["Municipal reconciliation"]
    for row in ws.iter_rows(min_row=5, max_row=9, values_only=True):
        if not row[0]:
            continue
        reg = row[0]
        gap = col[reg] - rcv[reg]
        bad = []
        for label, got, want in (("collections", row[1], int(round(col[reg]))),
                                 ("receipts", row[2], int(round(rcv[reg]))),
                                 ("gap", row[3], int(round(gap))),
                                 ("share", row[4], round(100.0 * gap / col[reg], 1))):
            if got != want:
                bad.append("%s %r != %r" % (label, got, want))
        ok("Municipal reconciliation row recomputes: %s" % reg, not bad, "; ".join(bad))

    # ---- the conventions the deliverables lean on, read back out of the documents that carry
    # them. Each absence or presence check runs through the same extractor and the same parse as
    # its own positive control, so a check that cannot read the file fails instead of passing.
    def flat(name):
        return " ".join(pdf_text(name).split())

    dsg = flat("bws_2026_design_and_weighting.pdf")
    ok("the design document states the field-day convention (control: the frame size)",
       "field days of eight hours" in dsg and "36,735 such enterprises" in dsg,
       "%d characters of text, convention present: %s"
       % (len(dsg), "field days of eight hours" in dsg))
    memo = flat("collection_capacity_note_2027.pdf")
    ok("the capacity memo states the municipal counting convention (control: the estimand clause)",
       "each municipality's collected tonnage counted once" in memo
       and "counted once at the producer's despatch leg" in memo,
       "%d characters of text" % len(memo))
    rel = flat("waste_statistics_2025_release.pdf")
    ok("the 2025 release states the extract its register column came from (control: the revision)",
       "as extracted on 30 June 2026" in rel and "supersedes the preliminary release" in rel,
       "%d characters of text" % len(rel))
    ok("no shipped document states hours per field day twice",
       sum(1 for nm in sorted(os.listdir(TGT)) if nm.endswith(".pdf")
           and "field days of eight hours" in flat(nm)) == 1,
       "the convention is filed once, in the survey's own collection documentation")

    # ---- Table 1 of the 2025 release against the register recompute. The release says the column
    # is an earlier extract and is not revised, so it has to sit visibly below the recompute: a
    # published column that nearly ties would read as a calibration target for the 2025 ask.
    lines = [l.strip() for l in pdf_text("waste_statistics_2025_release.pdf").splitlines()]
    pub = {}
    for i, l in enumerate(lines):
        if l.endswith("wastes") and l in name2key and i + 3 < len(lines):
            m = re.match(r"^([0-9]{1,2}\.[0-9]) per cent$", lines[i + 3])
            if m and lines[i + 1].isdigit() and lines[i + 2].isdigit():
                pub[name2key[l]] = (int(lines[i + 1]), int(lines[i + 2]), float(m.group(1)))
    ok("Table 1 of the 2025 release yields all eight rows", len(pub) == 8, "%d rows read" % len(pub))
    gaps = [(a2[k] - pub[k][1]) / float(a2[k]) for k in pub]
    ok("every published 2025 recovered tonnage sits below the register recompute by 2.5 to 6 per cent",
       all(0.025 <= g <= 0.06 for g in gaps),
       "gaps %s" % ", ".join("%.1f" % (100 * g) for g in sorted(gaps)))
    ok("every published 2025 rate is the published tonnages' own ratio",
       all(abs(round(100.0 * pub[k][1] / pub[k][0], 1) - pub[k][2]) < 0.05 for k in pub),
       "eight rows internally consistent")
    tot_txt = "{:,}".format(sum(v[0] for v in pub.values()))
    ok("the release's stated national total is the sum of its own arisings column",
       "estimated at %s tonnes" % tot_txt in rel, "total %s t" % tot_txt)

    ok("the chart ships", os.path.exists(os.path.join(GLD, "attainment_against_target_2026.png")),
       "attainment_against_target_2026.png")
    print("\n%s  verify: %d checks, %d FAIL" % ("PASS" if not FAILS else "FAIL", N[0], len(FAILS)))
    return 1 if FAILS else 0


if __name__ == "__main__":
    sys.exit(main())
