"""The two golden deliverables and figures.json. Every figure here is recomputed by verify.py from
the shipped bundle alone."""
import json
import math
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt                   # noqa: E402
from openpyxl import Workbook                     # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

import world as W                                 # noqa: E402

THIN = Side(style="thin", color="FF9E9E9E")
HDR = PatternFill("solid", fgColor="FFEDEDED")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _head(ws, row, labels, widths=None):
    for j, lab in enumerate(labels, start=1):
        c = ws.cell(row=row, column=j, value=lab)
        c.font = Font(bold=True, size=10)
        c.fill = HDR
        c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical="top")
    if widths:
        for col, wd in zip("ABCDEFGHIJKLMN", widths):
            ws.column_dimensions[col].width = wd
    ws.row_dimensions[row].height = 30


def workbook(path, att, chain, a1, a2, a3, prior):
    wb = Workbook()
    ws = wb.active
    ws.title = "Attainment"
    ws["A1"] = "Business waste recovery, attainment against the plan, reference year 2026"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Recovery rate: tonnage despatched in 2026 and confirmed received at a recovery "
                "operation, counted once at the producer's despatch leg, over estimated arisings.")
    ws["A3"] = ("A stream is established below target where the rate is below the plan target at "
                "the 95 per cent level. Standard error is of the rate; the register tonnage is a "
                "complete total and the arisings estimate is from the 2026 survey.")
    for r in (2, 3):
        ws.cell(row=r, column=1).alignment = Alignment(wrap_text=False)
    _head(ws, 5, ["Statistical stream", "Confirmed recovered (t)", "Estimated arisings (t)",
                  "Recovery rate (%)", "Standard error (pp)", "95% interval lower (%)",
                  "95% interval upper (%)", "Plan 2026 target (%)", "Shortfall (pp)",
                  "Standard errors from target", "Established below target"],
          [42, 15, 15, 12, 12, 13, 13, 13, 11, 13, 14])
    row = 6
    for k, name, _p, _m, _c in W.STREAMS:
        d = att[k]
        vals = [name, d["recovered_t"], int(round(d["arisings"])), round(d["rate"] * 100, 1),
                round(d["se"] * 100, 2), round(d["lo"] * 100, 1), round(d["hi"] * 100, 1),
                round(d["target"] * 100, 1), round(d["shortfall"] * 100, 1), round(d["z"], 2),
                "yes" if d["established"] else "no"]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=row, column=j, value=v)
            c.border = BOX
            if j > 1:
                c.alignment = Alignment(horizontal="right")
            if d["established"]:
                c.font = Font(bold=True)
        row += 1
    ws.cell(row=row + 1, column=1,
            value=("Module committed to: %s. It is the only stream whose shortfall is established; "
                   "%s carries the largest shortfall in points and is not established."
                   % (W.STREAM_NAME[W.GOLDEN_STREAM], W.STREAM_NAME[W.LOUD_STREAM])))
    _fmt(ws, 6, row - 1, {"B": "#,##0", "C": "#,##0", "D": "0.0", "E": "0.00", "F": "0.0",
                          "G": "0.0", "H": "0.0", "I": "0.0", "J": "0.00"})
    ws.freeze_panes = "B6"
    ws.auto_filter.ref = "A5:K%d" % (row - 1)
    ws.print_area = "A1:K%d" % (row + 1)
    ws.print_title_rows = "5:5"
    ws.page_setup.orientation = "landscape"

    ws = wb.create_sheet("Stage chain")
    ws["A1"] = "Tonnage by stage of the movement chain, 2026"
    ws["A1"].font = Font(bold=True, size=12)
    _head(ws, 3, ["Statistical stream", "1 Arisings, estimated (t)", "2 Despatched off site (t)",
                  "3 Confirmed at a final treatment operation (t)",
                  "4 Confirmed at a recovery operation (t)", "Pass through 1 to 2 (%)",
                  "Pass through 2 to 3 (%)", "Pass through 3 to 4 (%)"],
          [42, 16, 16, 20, 18, 13, 13, 13])
    row = 4
    for k, name, _p, _m, _c in W.STREAMS:
        c = chain[k]
        vals = [name, int(round(c["arisings"])), c["despatched"], c["confirmed"], c["recovered"],
                round(100.0 * c["despatched"] / c["arisings"], 1),
                round(100.0 * c["confirmed"] / c["despatched"], 1),
                round(100.0 * c["recovered"] / c["confirmed"], 1)]
        for j, v in enumerate(vals, start=1):
            cc = ws.cell(row=row, column=j, value=v)
            cc.border = BOX
            if j > 1:
                cc.alignment = Alignment(horizontal="right")
        row += 1
    _fmt(ws, 4, row - 1, {"B": "#,##0", "C": "#,##0", "D": "#,##0", "E": "#,##0",
                          "F": "0.0", "G": "0.0", "H": "0.0"})
    ws.freeze_panes = "B4"

    ws = wb.create_sheet("Collection effort")
    ws["A1"] = "Collection round 2026: contacts by region of the enterprise contacted"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("A case sits in the region of the enterprise it was opened for, which is the holder "
                "of its permit number over the round and not always the holder of record now. "
                "Contacts and usable returns are counted once.")
    _head(ws, 4, ["Region", "Contacts made", "Usable returns", "Contacts per usable return"],
          [18, 16, 14, 26])
    row = 5
    for r in W.REGIONS:
        for j, v in enumerate([r, a1[r]["contacts"], a1[r]["returns"],
                               round(a1[r]["per_return"], 2)], start=1):
            c = ws.cell(row=row, column=j, value=v)
            c.border = BOX
            if j > 1:
                c.alignment = Alignment(horizontal="right")
        row += 1
    _fmt(ws, 5, row - 1, {"B": "#,##0", "C": "#,##0", "D": "0.00"})
    ws.freeze_panes = "B5"

    ws = wb.create_sheet("Recovery 2025")
    ws["A1"] = "Tonnage confirmed received at a recovery operation, reference year 2025"
    ws["A1"].font = Font(bold=True, size=12)
    _head(ws, 3, ["Statistical stream", "Confirmed recovered 2025 (t)"], [42, 24])
    row = 4
    for k, name, _p, _m, _c in W.STREAMS:
        for j, v in enumerate([name, a2[k]], start=1):
            c = ws.cell(row=row, column=j, value=v)
            c.border = BOX
            if j > 1:
                c.alignment = Alignment(horizontal="right")
        row += 1
    _fmt(ws, 4, row - 1, {"B": "#,##0"})
    ws.freeze_panes = "B4"

    ws = wb.create_sheet("Municipal reconciliation")
    ws["A1"] = "Municipal waste 2026: reported collections against licensed facility receipts"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Receipts are on a receipt-date basis and attributed to the region of collection. "
                "A municipality's collections are its own return for the year, taken once.")
    _head(ws, 4, ["Region", "Reported collections (t)", "Facility receipts (t)", "Gap (t)",
                  "Gap as a share of collections (%)"], [18, 20, 18, 14, 22])
    row = 5
    for r in W.REGIONS:
        g = a3[r]
        for j, v in enumerate([r, int(round(g["collections"])), int(round(g["receipts"])),
                               int(round(g["gap"])), round(g["share"], 1)], start=1):
            c = ws.cell(row=row, column=j, value=v)
            c.border = BOX
            if j > 1:
                c.alignment = Alignment(horizontal="right")
        row += 1
    _fmt(ws, 5, row - 1, {"B": "#,##0", "C": "#,##0", "D": "#,##0", "E": "0.0"})
    ws.freeze_panes = "B5"

    ws = wb.create_sheet("Notes")
    notes = [
        ("Reference year", "2026, on the producer's despatch date"),
        ("Register extract", W.EXTRACT_REGISTER),
        ("Survey round", "%s, 2026 round" % W.EXTRACT_SURVEY),
        ("Recovery rate", "tonnage despatched in the reference year and confirmed received at a "
                          "recovery operation, counted once at the producer's despatch leg, over "
                          "estimated arisings"),
        ("Arisings", "total tonnage arising at every enterprise in the survey frame, sampled or "
                     "not, estimated from the 2026 round"),
        ("Excluded", "consignments despatched in the reference year and still unconfirmed at the "
                     "extract date; these are shown in the stage chain as the fall from stage 2 to "
                     "stage 3"),
        ("Interim operations", "a holding site's confirmation under an interim code is not a "
                               "recovery operation; the operation confirmed on the consignment the "
                               "holding site forwards is the one that counts"),
        ("Collection case", "one per enterprise issued, in the region of the enterprise the case "
                            "was opened for; the response log carries the permit number the case "
                            "was opened under and the crosswalk is effective dated"),
        ("Municipal collections", "each municipality's own return, counted once; the "
                                  "inter-municipal associations report for their members"),
        ("Rounding", "rates, shortfalls and shares to one decimal, tonnages to the nearest tonne, "
                     "contacts to the nearest whole and contacts per usable return to two decimals"),
        ("Standard error", "of the recovery rate, shown for working; the plan reports attainment "
                           "without a precision measure"),
        ("Source", "%s and the %s" % (W.SURVEY, W.REGISTER)),
    ]
    ws["A1"] = "Definitions and conventions"
    ws["A1"].font = Font(bold=True, size=12)
    r = 3
    for k, v in notes:
        a = ws.cell(row=r, column=1, value=k)
        a.font = Font(bold=True, size=9)
        a.alignment = Alignment(vertical="top")
        b = ws.cell(row=r, column=2, value=v)
        b.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 14 if len(v) < 70 else 28
        r += 1
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 92
    wb.save(path)


def _fmt(ws, r0, r1, spec):
    for col, nf in spec.items():
        for r in range(r0, r1 + 1):
            ws["%s%d" % (col, r)].number_format = nf


def chart(path, att):
    keys = [k for k, _n, _p, _m, _c in W.STREAMS]
    order = sorted(keys, key=lambda k: att[k]["z"])
    names = [W.STREAM_NAME[k] for k in order]
    rate = [att[k]["rate"] * 100 for k in order]
    err = [att[k]["se"] * W.BAR * 100 for k in order]
    tgt = [att[k]["target"] * 100 for k in order]
    est = [att[k]["established"] for k in order]
    fig, ax = plt.subplots(figsize=(9.2, 5.0), dpi=150)
    y = range(len(order))
    bars = ax.barh(list(y), rate, height=0.56,
                   color=["#9fb7c9" if not e else "#1f4e66" for e in est],
                   edgecolor="#5a6b78", linewidth=0.5, zorder=2)
    ax.errorbar(rate, list(y), xerr=err, fmt="none", ecolor="#2b2b2b", elinewidth=1.1,
                capsize=3.5, zorder=3)
    for i, t in enumerate(tgt):
        ax.plot([t, t], [i - 0.33, i + 0.33], color="#b4472a", linewidth=2.0, zorder=4)
    ax.set_yticks(list(y))
    ax.set_yticklabels(names, fontsize=8.5)
    ax.set_xlabel("2026 recovery rate, per cent (bar), 95 per cent interval (whisker), "
                  "plan 2026 target (vertical mark)", fontsize=8.5)
    ax.set_title("Business waste recovery against the plan, 2026: one stream is established below "
                 "target", fontsize=10.5, loc="left")
    ax.set_xlim(0, 78)
    ax.grid(axis="x", color="#dddddd", linewidth=0.6, zorder=1)
    ax.set_axisbelow(True)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    for i, e in enumerate(est):
        if e:
            ax.annotate("established below target", (max(rate[i] + err[i], tgt[i]) + 1.8, i),
                        fontsize=8, va="center", color="#1f4e66")
            ax.annotate("target %.1f" % tgt[i], (tgt[i], i + 0.42), fontsize=7.2, ha="center",
                        va="bottom", color="#b4472a")
    fig.text(0.008, 0.012, "Source: %s, 2026 round, and the %s, extract of %s. Plan targets: %s, "
             "Annex II Table 4." % (W.SURVEY, W.REGISTER, W.EXTRACT_REGISTER, W.PLAN),
             fontsize=6.6, color="#555555")
    fig.tight_layout(rect=(0, 0.035, 1, 1))
    fig.savefig(path, format="png", metadata={"Software": None})
    plt.close(fig)


def figures(path, att, chain, a1, a2, a3, extra):
    g = att[W.GOLDEN_STREAM]
    out = {
        "call": W.STREAM_NAME[W.GOLDEN_STREAM],
        "call_rate_pct": round(g["rate"] * 100, 1),
        "call_target_pct": round(g["target"] * 100, 1),
        "call_shortfall_pp": round(g["shortfall"] * 100, 1),
        "call_se_pp": round(g["se"] * 100, 2),
        "call_z": round(g["z"], 2),
        "call_interval_pct": [round(g["lo"] * 100, 1), round(g["hi"] * 100, 1)],
        "runner_up": W.STREAM_NAME[W.LOUD_STREAM],
        "runner_up_rate_pct": round(att[W.LOUD_STREAM]["rate"] * 100, 1),
        "runner_up_shortfall_pp": round(att[W.LOUD_STREAM]["shortfall"] * 100, 1),
        "runner_up_z": round(att[W.LOUD_STREAM]["z"], 2),
        "established_streams": [W.STREAM_NAME[k] for k, *_ in W.STREAMS if att[k]["established"]],
    }
    for k, name, _p, _m, _c in W.STREAMS:
        d = att[k]
        out["stream_%s" % k] = {
            "recovered_t": d["recovered_t"], "arisings_t": int(round(d["arisings"])),
            "rate_pct": round(d["rate"] * 100, 1), "se_pp": round(d["se"] * 100, 2),
            "interval_pct": [round(d["lo"] * 100, 1), round(d["hi"] * 100, 1)],
            "target_pct": round(d["target"] * 100, 1),
            "shortfall_pp": round(d["shortfall"] * 100, 1), "z": round(d["z"], 2),
            "established": d["established"]}
        c = chain[k]
        out["chain_%s" % k] = {
            "arisings_t": int(round(c["arisings"])), "despatched_t": c["despatched"],
            "confirmed_t": c["confirmed"], "recovered_t": c["recovered"],
            "pass_1_2_pct": round(100.0 * c["despatched"] / c["arisings"], 1),
            "pass_2_3_pct": round(100.0 * c["confirmed"] / c["despatched"], 1),
            "pass_3_4_pct": round(100.0 * c["recovered"] / c["confirmed"], 1)}
    for r in W.REGIONS:
        out["a1_%s" % r] = {"contacts": a1[r]["contacts"],
                            "usable_returns": a1[r]["returns"],
                            "contacts_per_return": round(a1[r]["per_return"], 2)}
        out["a3_%s" % r] = {"collections_t": int(round(a3[r]["collections"])),
                            "receipts_t": int(round(a3[r]["receipts"])),
                            "gap_t": int(round(a3[r]["gap"])), "gap_share_pct": round(a3[r]["share"], 1)}
    for k, name, _p, _m, _c in W.STREAMS:
        out["a2_%s" % k] = a2[k]
    out.update(extra)
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(out, fh, indent=2, sort_keys=True)
    return out
