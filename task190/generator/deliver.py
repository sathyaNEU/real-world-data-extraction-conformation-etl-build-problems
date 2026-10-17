"""The two files the reporting officer puts in front of the committee."""
import os

import world as W


def workbook(path, g):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "7.4 structuring"
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"] = "%s: cash deposits split below the prescribed transaction threshold" % W.BANK
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Paragraph 7.4 of the programme handbook, reporting year ended 30 June 2027. "
                "Counts of deposits to the nearest 10.")
    ws["A2"].font = Font(size=9, italic=True)
    ws["A4"] = "Deposits split"
    ws["A4"].font = Font(bold=True)
    ws["B4"] = g["total"]
    ws["B4"].font = Font(bold=True)
    ws["A5"] = "Largest segment"
    ws["A5"].font = Font(bold=True)
    ws["B5"] = "%s, %d" % (W.SEG_NAME[g["largest"]], g["per_seg"][g["largest"]])
    head = ["Business segment", "Cash deposits paid in", "Deposits split"]
    for j, h in enumerate(head):
        c = ws.cell(row=7, column=1 + j, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8E8E4")
        c.border = box
        c.alignment = Alignment(horizontal="center" if j else "left")
    order = sorted(W.SEG_KEYS, key=lambda k: (-g["per_seg"][k], W.SEG_NAME[k]))
    r = 8
    for k in order:
        for j, v in enumerate([W.SEG_NAME[k], g["deposits"][k], g["per_seg"][k]]):
            c = ws.cell(row=r, column=1 + j, value=v)
            c.border = box
            if j:
                c.alignment = Alignment(horizontal="right")
                c.number_format = "#,##0"
        r += 1
    for j, v in enumerate(["Total", sum(g["deposits"].values()), g["total"]]):
        c = ws.cell(row=r, column=1 + j, value=v)
        c.font = Font(bold=True)
        c.border = box
        if j:
            c.alignment = Alignment(horizontal="right")
            c.number_format = "#,##0"
    ws.cell(row=r + 2, column=1,
            value="Cash deposits paid in counts every physical cash deposit of the segment's "
                  "customers in the reporting year, at any amount.").font = Font(size=9, italic=True)
    ws.cell(row=r + 3, column=1,
            value="Teller and machine feeds pulled 19 July 2027. Financial crime team, for the "
                  "audit and risk committee.").font = Font(size=9, italic=True)
    ws.cell(row=r + 4, column=1, value="Notes").font = Font(bold=True)
    for i, line in enumerate([
            "Deposits split is the paragraph 7.4 count, rounded to the nearest 10 deposits; the "
            "total is the sum of the segment rows.",
            "Segments are the eight the customer register carries. Codes were reissued on 6 April "
            "2027 and the segments were not changed (handbook 7.5).",
            "Cash is notes and coin at a counter or a smart deposit machine; cheques and electronic "
            "credits are outside the count (handbook 3.1).",
            "The Hastings Heretaunga and Hastings Karamu branches are treated as one for the whole "
            "year (handbook 9.3).",
            "Two reversed deposits and their reversing rows are excluded."]):
        ws.cell(row=r + 5 + i, column=1, value=line).font = Font(size=9)
    for col, wd in (("A", 34), ("B", 24), ("C", 18)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A8"
    ws.print_area = "A1:C%d" % (r + 10)
    ws.print_title_rows = "7:7"
    wb.save(path)
    return path


def chart(path, g):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    order = sorted(W.SEG_KEYS, key=lambda k: (-g["per_seg"][k], W.SEG_NAME[k]))
    vals = [g["per_seg"][k] for k in order]
    names = [W.SEG_NAME[k].replace(" and ", " &\n") for k in order]
    fig, ax = plt.subplots(figsize=(8.8, 4.6), dpi=150)
    cols = ["#1f4e79" if k == g["largest"] else "#8099b4" for k in order]
    ax.bar(names, vals, color=cols, edgecolor="#35485c", linewidth=0.6)
    for i, v in enumerate(vals):
        ax.annotate("%d" % v, xy=(i, v), xytext=(0, 3), textcoords="offset points",
                    ha="center", fontsize=9)
    ax.set_ylabel("Deposits split, to the nearest 10")
    ax.set_title("%s carries the most deposits split below the threshold"
                 % W.SEG_NAME[g["largest"]], loc="left", fontsize=12, fontweight="bold",
                 pad=24)
    ax.text(0.0, 1.015, "Paragraph 7.4 count by business segment, year ended 30 June 2027",
            transform=ax.transAxes, fontsize=9, color="#55606b")
    fig.text(0.012, 0.012, "Source: teller and smart machine deposit feeds pulled 19 July 2027, "
                           "business banking register. Counts to the nearest 10 deposits.",
             fontsize=7.6, color="#55606b")
    ax.set_ylim(0, max(vals) * 1.22)
    ax.yaxis.grid(True, color="#e2e2de", linewidth=0.7)
    ax.set_axisbelow(True)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(axis="x", labelsize=8)
    fig.tight_layout(rect=(0, 0.04, 1, 1))
    fig.savefig(path, metadata={"Software": None, "Creation Time": None})
    plt.close(fig)
    return path
