"""The golden, computed from the shipped bundle only, and every route the checks hold it against.

The golden: each casual's answer to a call-in offer is a binary choice, fitted on every offer in the history with the
offer's position against the casual's own roster that day as a term (an offer that starts as the casual's rostered or
booked shift at that store ends, or ends as it starts, needs no separate trip to work); the fitted choice model is then
run through the call-in module as its notes describe, over the Christmas 2026 rosters as published, with every
store's open shifts posted at 09:00 on Monday 16 November, and each store's open paid hours left unfilled are averaged
over the runs. The crew goes to the store with the most; its crew hours are that store's open paid hours."""
from __future__ import annotations

import bisect
import collections
import csv
import datetime as dt
import heapq
import json
import math
import os

import numpy as np
from openpyxl import load_workbook

import common as K

# ------------------------------------------------------------------------------------------------ calendar and clock
D0 = dt.date(2025, 8, 29)
XMAS25 = (dt.date(2025, 11, 28), dt.date(2025, 12, 24))
XMAS26 = (dt.date(2026, 11, 27), dt.date(2026, 12, 24))
POST26 = dt.datetime(2026, 11, 16, 9, 0)


def dn(d):
    return (d - D0).days


def utc_to_local_min(s, convert=True):
    """'2025-12-04T20:02:13Z' -> minutes from D0 00:00 store time (convert=False reads the UTC clock as store time)"""
    t = dt.datetime.strptime(s, "%Y-%m-%dT%H:%M:%SZ")
    if convert:
        t = t.replace(tzinfo=K.UTC).astimezone(K.MEL).replace(tzinfo=None)
    return dn(t.date()) * 1440 + t.hour * 60 + t.minute + t.second / 60.0


def hm(s):
    h, m = s.split(":")
    return int(h) * 60 + int(m)


def ampm(s):
    tm, suf = s.split(" ")
    h, m = tm.split(":")
    h = int(h) % 12 + (12 if suf == "PM" else 0)
    return h * 60 + int(m)


def paid_min(span):
    return span - (30 if span > 300 else 0)


def in_x25(d):
    return XMAS25[0] <= d <= XMAS25[1]


def is_xmas(d):
    return in_x25(d) or XMAS26[0] <= d <= XMAS26[1]


# ------------------------------------------------------------------------------------------------ the bundle
def read_csv(path):
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


class Bundle:
    pass


def load(target):
    B = Bundle()
    B.target = target
    B.stores = read_csv(os.path.join(target, K.F_STORES))
    B.loc_store = {r["kws_location"]: r["store_name"] for r in B.stores}
    B.cc_loc = {r["cost_centre"]: r["kws_location"] for r in B.stores}
    B.name_loc = {r["store_name"].upper(): r["kws_location"] for r in B.stores}
    B.cluster = {r["kws_location"]: r["callin_cluster"] for r in B.stores}
    B.locs = [r["kws_location"] for r in B.stores]
    B.people = {r["person_code"]: r for r in read_csv(os.path.join(target, K.F_PEOPLE))}
    B.payroll = read_csv(os.path.join(target, K.F_PAYROLL))
    B.offers = read_csv(os.path.join(target, K.F_OFFERS))
    B.register = read_csv(os.path.join(target, K.F_REGISTER))
    B.avail = json.load(open(os.path.join(target, K.F_AVAIL), encoding="utf-8"))
    B.roster = read_rosters(B)
    B.xmas = read_christmas(B)
    identities(B)
    return B


def read_rosters(B):
    """every published roster line of the history: (loc, day, start, end, payroll or None, status)"""
    out = []
    for r in read_csv(os.path.join(B.target, K.F_LEGACY)):
        d = dt.datetime.strptime(r["Date"], "%d/%m/%Y").date()
        s, e = ampm(r["Start"]), ampm(r["Finish"])
        if e <= s:
            e += 1440
        pn = int(r["Emp No"]) if r["Emp No"] else None
        out.append((B.name_loc[r["Store"]], dn(d), s, e, pn, r["Status"] or "ROSTERED"))
    for r in read_csv(os.path.join(B.target, K.F_KWS)):
        d = dt.date.fromisoformat(r["shift_date"])
        s, e = hm(r["start_time"]), hm(r["end_time"])
        pn = int(r["payroll_no"]) if r["payroll_no"] else None
        out.append((r["location_id"], dn(d), s, e, pn, r["line_status"]))
    return out


def read_christmas(B):
    wb = load_workbook(os.path.join(B.target, K.F_XMAS), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        loc = ws["A2"].value.split("(")[1].rstrip(")")
        for row in ws.iter_rows(min_row=7, values_only=True):
            if row[0] is None:
                continue
            d = row[0].date() if isinstance(row[0], dt.datetime) else row[0]
            s, e = hm(row[3]), hm(row[4])
            out.append(dict(loc=loc, day=dn(d), s=s, e=e, code=row[2], pn=int(row[6]) if row[6] is not None else None,
                            typ=row[7] or "", open=row[8] == "OPEN"))
    return out


def identities(B, follow_rehires=True):
    """person code <-> payroll numbers through rehires; home location on a date from payroll's cost-centre history"""
    B.eng = collections.defaultdict(list)          # employee_no -> [(from, to, loc)], inside the engagement
    B.cc_raw = collections.defaultdict(list)       # employee_no -> [(from, to, loc)] as the rows give them
    B.prev = {}
    B.dob = {}
    for r in B.payroll:
        no = int(r["employee_no"])
        f = dt.date.fromisoformat(r["cc_effective_from"])
        t = dt.date.fromisoformat(r["cc_effective_to"]) if r["cc_effective_to"] else dt.date(2099, 1, 1)
        a = dt.date.fromisoformat(r["engagement_start"])
        b = dt.date.fromisoformat(r["engagement_end"]) if r["engagement_end"] else dt.date(2099, 1, 1)
        loc = B.cc_loc[r["cost_centre"]]
        B.cc_raw[no].append((f, t, loc))
        # a cost centre applies to the engagement it belongs to, never before the engagement starts
        if max(f, a) <= min(t, b):
            B.eng[no].append((max(f, a), min(t, b), loc))
        B.dob[no] = dt.date.fromisoformat(r["date_of_birth"])
        if r["previous_employee_no"]:
            B.prev[no] = int(r["previous_employee_no"])
    B.nos = {}                                       # person_code -> [payroll numbers, oldest first]
    B.person_of = {}
    for pc, r in B.people.items():
        chain = [int(r["payroll_no"])]
        while follow_rehires and chain[0] in B.prev:
            chain.insert(0, B.prev[chain[0]])
        B.nos[pc] = chain
        for no in chain:
            B.person_of[no] = pc


def home_loc(B, pc, d):
    for no in B.nos[pc]:
        for f, t, loc in B.eng[no]:
            if f <= d <= t:
                return loc
    return B.people[pc]["home_location"]


def home_loc_newest(B, pc, d):
    """the home store read from the number KWS holds first, on the cost-centre rows' own dates, earlier numbers only
    when it has no row for the date (the natural read, which takes a re-engagement's cost centre from the day it was
    keyed)"""
    for no in reversed(B.nos[pc]):
        for f, t, loc in B.cc_raw[no]:
            if f <= d <= t:
                return loc
    return B.people[pc]["home_location"]


def payroll_on(B, pc, d):
    for no in B.nos[pc]:
        for f, t, loc in B.eng[no]:
            if f <= d <= t:
                return no
    return B.nos[pc][-1]


def age_on(B, pc, d):
    dob = B.dob[B.nos[pc][-1]]
    return (d - dob).days / 365.25


# ------------------------------------------------------------------------------------------------ history features
def history(B, convert=True, once=True):
    """the offer log with each offer's terms, built from the rosters and the call-in register"""
    vac = set()
    filled = collections.defaultdict(list)      # (person, day) -> [(loc, s, e, t_acc, t_drop)]
    regs = {}
    for r in B.register:
        s0 = utc_to_local_min(r["shift_start_utc"])
        e0 = utc_to_local_min(r["shift_end_utc"])
        day = int(s0 // 1440)
        s, e = s0 - day * 1440, e0 - day * 1440
        regs[r["callin_id"]] = (r["location_id"], day, int(round(s)), int(round(e)))
        if r["reason"] in ("ABS", "LVE") and r["absent_payroll_no"]:
            vac.add((r["location_id"], day, int(round(s)), int(round(e)), int(r["absent_payroll_no"])))
    drops = {}
    for r in sorted(B.register, key=lambda r: r["raised_at"]):
        if r["reason"] == "ABS" and r["absent_payroll_no"]:
            no = int(r["absent_payroll_no"])
            pc = B.person_of.get(no)
            if pc:
                loc, day, s, e = regs[r["callin_id"]]
                drops.setdefault((pc, loc, day, s, e), utc_to_local_min(r["raised_at"]))
    for r in B.register:
        if r["outcome"] == "FILLED":
            loc, day, s, e = regs[r["callin_id"]]
            t = utc_to_local_min(r["closed_at"])
            td = drops.get((r["filled_by"], loc, day, s, e))
            item = (loc, s, e, t, td if td is not None and td > t else None)
            # a casual is booked on a shift once, however many call-ins for it closed FILLED with them
            if not once or item not in filled[(r["filled_by"], day)]:
                filled[(r["filled_by"], day)].append(item)
    book = collections.defaultdict(list)        # (payroll, day) -> [(loc, s, e)]
    for loc, day, s, e, pn, st in B.roster:
        if pn is not None and st == "ROSTERED" and (loc, day, s, e, pn) not in vac:
            book[(pn, day)].append((loc, s, e))
    rows = []
    asked = set()
    for o in B.offers:
        pc = o["person_code"]
        d = dt.date.fromisoformat(o["shift_date"])
        day = dn(d)
        s, e = hm(o["shift_start"]), hm(o["shift_end"])
        t = utc_to_local_min(o["offered_at"], convert)
        loc = o["location_id"]
        today = []
        pn = payroll_on(B, pc, d)
        today += book.get((pn, day), [])
        for (l2, s2, e2, ta, td) in filled.get((pc, day), []):
            if ta <= t and (td is None or td > t) and not (l2 == loc and s2 == s and e2 == e):
                today.append((l2, s2, e2))
        ext = any(l2 == loc and (e2 == s or s2 == e) for l2, s2, e2 in today)
        split = (not ext) and bool(today)
        # the module notes' travel rule, read back on the record: a shift within an hour of one at another store
        near = any(l2 != loc and ((e2 <= s and s - e2 < TRAVEL_MIN) or (s2 >= e and s2 - e < TRAVEL_MIN))
                   for l2, s2, e2 in today)
        wk = day // 7
        hrs = 0.0
        for k in range(wk * 7, wk * 7 + 7):
            for l2, s2, e2 in book.get((pn, k), []):
                hrs += paid_min(e2 - s2) / 60.0
            for (l2, s2, e2, ta, td) in filled.get((pc, k), []):
                if ta <= t and (td is None or td > t):
                    hrs += paid_min(e2 - s2) / 60.0
        key = (pc, loc, day, s, e)
        re_ = key in asked
        asked.add(key)
        home = home_loc(B, pc, d)
        rows.append(dict(y=1 if o["response"] == "ACCEPTED" else 0, ext=ext, split=split, near=near, cl=home != loc,
                         hours_wk=hrs, len_h=paid_min(e - s) / 60.0, late=e > 1290, early=s < 420,
                         sun=d.weekday() == 6, short_notice=(day * 1440 + s - t) < 240, reoffer=re_, dec=in_x25(d),
                         home=home, loc=loc, age=age_on(B, pc, d), pc=pc, day=day, answer=o["response"]))
    return rows


# ------------------------------------------------------------------------------------------------ the choice model
SPECS = {
    # the golden: the trip-cost term with the season, the casual's age band and home store
    "golden": dict(ext=True, age=True, store=True),
    "ext_only": dict(ext=True),
    "ext_age": dict(ext=True, age=True),
    "ext_dec": dict(ext=True, age=True, store=True, extdec=True),
    "noext": dict(),
    "noext_age": dict(age=True),
    "noext_age_store": dict(age=True, store=True),
    "sameday": dict(age=True, same=True),
    "sameday_store": dict(age=True, store=True, same=True),
}


def design(rows, spec, locs):
    X = []
    for f in rows:
        x = [1.0, f["len_h"], f["cl"], max(0.0, f["hours_wk"] - 10.0), f["late"], f["early"], f["sun"],
             f["short_notice"], f["reoffer"], f["dec"]]
        if spec.get("age"):
            a = f["age"]
            x += [a < 18, 18 <= a < 25]
        if spec.get("store"):
            x += [f["home"] == l for l in locs[1:]]
        if spec.get("same"):
            x.append(f["ext"] or f["split"])
        if spec.get("ext"):
            x += [f["ext"], f["split"]]
        if spec.get("extdec"):
            x.append(f["ext"] and f["dec"])
        X.append(x)
    return np.array(X, dtype=float)


def fit_logit(X, y, iters=60):
    """Newton-Raphson maximum likelihood with a tiny ridge for separation"""
    b = np.zeros(X.shape[1])
    for _ in range(iters):
        z = X @ b
        p = 1.0 / (1.0 + np.exp(-z))
        g = X.T @ (y - p) - 1e-6 * b
        H = (X * (p * (1 - p))[:, None]).T @ X + 1e-6 * np.eye(X.shape[1])
        step = np.linalg.solve(H, g)
        b += step
        if np.max(np.abs(step)) < 1e-9:
            break
    cov = np.linalg.inv(H)
    return b, cov


def fit(rows, spec_name, locs, mask=None):
    spec = SPECS[spec_name]
    rs = rows if mask is None else [r for r, m in zip(rows, mask) if m]
    X = design(rs, spec, locs)
    y = np.array([r["y"] for r in rs], dtype=float)
    b, cov = fit_logit(X, y)
    return dict(spec=spec_name, b=b, cov=cov, n=len(rs))


def prob(model, f, locs):
    x = design([f], SPECS[model["spec"]], locs)[0]
    return 1.0 / (1.0 + math.exp(-float(x @ model["b"])))


# ------------------------------------------------------------------------------------------------ the forward window
TRAVEL_MIN = 60     # the module notes: no offer leaves a casual under an hour to get between stores
DAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


def windows_on(B, pc, d):
    """the casual's declared windows on date d from their availability profile"""
    prof = B.avail_idx.get(pc)
    if not prof:
        return []
    k = bisect.bisect_right(prof[0], d) - 1
    if k < 0:
        return []
    f, t, weekly = prof[1][k]
    if t is not None and d > t:
        return []
    return weekly[DAYS[d.weekday()]]


def index_availability(B):
    B.avail_idx = {}
    for p in B.avail["people"]:
        starts, vers = [], []
        for v in p["profiles"]:
            f = dt.date.fromisoformat(v["effective_from"])
            t = dt.date.fromisoformat(v["effective_to"]) if v["effective_to"] else None
            weekly = {k: [(hm(a), hm(b)) for a, b in w] for k, w in v["weekly"].items()}
            starts.append(f)
            vers.append((f, t, weekly))
        B.avail_idx[p["person_code"]] = (starts, vers)


class Forward:
    """the call-in module over the Christmas 2026 rosters as published, with a fitted choice model"""

    def __init__(self, B, model, home_rule="kws"):
        self.B = B
        if not hasattr(B, "avail_idx"):
            index_availability(B)
        self.model = model
        self.locs = B.locs
        d_as = XMAS26[0]
        act = [pc for pc, r in B.people.items() if r["status"] == "ACTIVE"]
        self.home = {pc: B.people[pc]["home_location"] for pc in act}
        self.age = {pc: age_on(B, pc, dt.date(2026, 12, 10)) for pc in act}
        svc = {pc: B.people[pc]["service_date"] for pc in act}
        self.order = {}
        for loc in self.locs:
            cl = B.cluster[loc]
            own = sorted([pc for pc in act if self.home[pc] == loc], key=lambda pc: (svc[pc], pc))
            oth = sorted([pc for pc in act if self.home[pc] != loc and cl and B.people[pc]["cluster_callins"] == "Y"
                          and B.cluster[self.home[pc]] == cl], key=lambda pc: (svc[pc], pc))
            self.order[loc] = own + oth
        pn_pc = {int(B.people[pc]["payroll_no"]): pc for pc in act}
        self.base = collections.defaultdict(list)
        self.open = []
        for ln in B.xmas:
            if ln["open"]:
                self.open.append(ln)
            elif ln["pn"] in pn_pc:
                self.base[(pn_pc[ln["pn"]], ln["day"])].append((ln["loc"], ln["s"], ln["e"]))
        self.post = dn(POST26.date()) * 1440 + POST26.hour * 60

    def reset(self):
        self.book = collections.defaultdict(list, {k: list(v) for k, v in self.base.items()})
        self.wk = collections.defaultdict(float)
        self.dm = collections.defaultdict(float)
        for (pc, day), lst in self.book.items():
            for loc, s, e in lst:
                self.wk[(pc, day // 7)] += paid_min(e - s) / 60.0
                self.dm[(pc, day)] += paid_min(e - s)

    def eligible(self, pc, ln):
        day, s, e = ln["day"], ln["s"], ln["e"]
        d = D0 + dt.timedelta(days=day)
        if not any(a <= s and e <= b for a, b in windows_on(self.B, pc, d)):
            return False
        for l2, s2, e2 in self.book.get((pc, day), ()):
            if s2 < e and s < e2:
                return False
            if l2 != ln["loc"] and ((e2 <= s and s - e2 < TRAVEL_MIN) or (s2 >= e and s2 - e < TRAVEL_MIN)):
                return False
        for l2, s2, e2 in self.book.get((pc, day - 1), ()):
            if (day * 1440 + s) - ((day - 1) * 1440 + e2) < 600:
                return False
        for l2, s2, e2 in self.book.get((pc, day + 1), ()):
            if ((day + 1) * 1440 + s2) - (day * 1440 + e) < 600:
                return False
        ln_ = paid_min(e - s)
        if self.dm[(pc, day)] + ln_ > 660 or self.wk[(pc, day // 7)] + ln_ / 60.0 > 38.0:
            return False
        return True

    def features(self, pc, ln, t, reoffer):
        day, s, e = ln["day"], ln["s"], ln["e"]
        today = self.book.get((pc, day), ())
        ext = any(l2 == ln["loc"] and (e2 == s or s2 == e) for l2, s2, e2 in today)
        d = D0 + dt.timedelta(days=day)
        return dict(ext=ext, split=(not ext) and bool(today), cl=self.home[pc] != ln["loc"],
                    hours_wk=self.wk[(pc, day // 7)], len_h=paid_min(e - s) / 60.0, late=e > 1290, early=s < 420,
                    sun=d.weekday() == 6, short_notice=(day * 1440 + s - t) < 240, reoffer=reoffer, dec=True,
                    home=self.home[pc], age=self.age[pc])

    def static(self):
        """each open shift's chance that every eligible casual says no, at posting, independently"""
        self.reset()
        u = collections.Counter()
        for ln in self.open:
            lq = 0.0
            for pc in self.order[ln["loc"]]:
                if self.eligible(pc, ln):
                    lq += math.log(max(1e-12, 1 - prob(self.model, self.features(pc, ln, self.post, False), self.locs)))
            u[ln["loc"]] += paid_min(ln["e"] - ln["s"]) / 60.0 * math.exp(lq)
        return u

    def simulate(self, runs=12, seed=0):
        """the module's cascade: one offer at a time down the list, 07:00 to 21:00, a decline final, the list
        offered again at 07:00 the next day; answers drawn from the fitted model; mean unfilled hours by store"""
        tot = collections.Counter()
        for r in range(runs):
            rng = np.random.default_rng([seed, r])
            self.reset()
            heap = []
            st = []
            for k, ln in enumerate(sorted(self.open, key=lambda l: (l["day"], l["s"], l["loc"], l["e"]))):
                st.append(dict(ln=ln, ptr=0, declined=set(), asked=set(), done=False))
                heapq.heappush(heap, (self.post, k))
            while heap:
                t, k = heapq.heappop(heap)
                x = st[k]
                if x["done"]:
                    continue
                ln = x["ln"]
                start = ln["day"] * 1440 + ln["s"]
                if t >= start - 15:
                    x["done"] = True
                    tot[ln["loc"]] += paid_min(ln["e"] - ln["s"]) / 60.0
                    continue
                lt = t % 1440
                if start - t > 720 and not (420 <= lt < 1260):
                    heapq.heappush(heap, ((t // 1440) * 1440 + 420 + (1440 if lt >= 1260 else 0), k))
                    continue
                lst = self.order[ln["loc"]]
                found = None
                while x["ptr"] < len(lst):
                    pc = lst[x["ptr"]]
                    if pc not in x["declined"] and self.eligible(pc, ln):
                        found = pc
                        break
                    x["ptr"] += 1
                if found is None:
                    nxt = (t // 1440) * 1440 + 420
                    if nxt <= t:
                        nxt += 1440
                    if nxt >= start - 30:
                        x["done"] = True
                        tot[ln["loc"]] += paid_min(ln["e"] - ln["s"]) / 60.0
                        continue
                    x["ptr"] = 0
                    heapq.heappush(heap, (nxt, k))
                    continue
                p = prob(self.model, self.features(found, ln, t, found in x["asked"]), self.locs)
                x["asked"].add(found)
                if rng.random() < p:
                    day = ln["day"]
                    self.book[(found, day)].append((ln["loc"], ln["s"], ln["e"]))
                    self.wk[(found, day // 7)] += paid_min(ln["e"] - ln["s"]) / 60.0
                    self.dm[(found, day)] += paid_min(ln["e"] - ln["s"])
                    x["done"] = True
                    continue
                x["declined"].add(found)
                x["ptr"] += 1
                heapq.heappush(heap, (t + 14.0, k))
        return {loc: tot[loc] / runs for loc in self.locs}


def open_hours(B):
    u = collections.Counter()
    for ln in B.xmas:
        if ln["open"]:
            u[ln["loc"]] += paid_min(ln["e"] - ln["s"]) / 60.0
    return {loc: u[loc] for loc in B.locs}


def rank(u):
    order = sorted(u, key=lambda l: -u[l])
    return order, (u[order[0]] / u[order[1]] if u[order[1]] > 0 else float("inf"))


# ------------------------------------------------------------------------------------------------ register routes
def register_shares(B, window):
    """each store's share of call-in hours filled, over the register rows whose shift falls in `window`"""
    tot = collections.Counter()
    fil = collections.Counter()
    for r in B.register:
        s0 = utc_to_local_min(r["shift_start_utc"])
        e0 = utc_to_local_min(r["shift_end_utc"])
        d = D0 + dt.timedelta(days=int(s0 // 1440))
        if not window(d):
            continue
        h = paid_min(e0 - s0) / 60.0
        tot[r["location_id"]] += h
        if r["outcome"] == "FILLED":
            fil[r["location_id"]] += h
    return {loc: fil[loc] / tot[loc] if tot[loc] else 0.0 for loc in B.locs}


def golden(B, runs=24):
    rows = history(B)
    model = fit(rows, "golden", B.locs)
    fw = Forward(B, model)
    u = fw.simulate(runs=runs, seed=155)
    oh = open_hours(B)
    order, margin = rank(u)
    return dict(rows=rows, model=model, unworked=u, open=oh, order=order, margin=margin, store=order[0],
                crew_hours=oh[order[0]])
