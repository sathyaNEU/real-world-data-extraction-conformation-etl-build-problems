"""The goldens: Workforce Planning's one-page placement brief (PDF, with the chart) and its workbook, both drawn from
the golden computed on the shipped bundle, and generator/figures.json with every figure submission.md states."""
from __future__ import annotations

import collections
import datetime as dt
import io
import json
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt                                              # noqa: E402
import matplotlib.ticker                                                     # noqa: E402
from openpyxl import Workbook                                                # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill                     # noqa: E402
from openpyxl.utils import get_column_letter                                 # noqa: E402
from reportlab import rl_config                                              # noqa: E402

rl_config.invariant = 1
from reportlab.lib import colors                                             # noqa: E402
from reportlab.lib.pagesizes import A4                                       # noqa: E402
from reportlab.lib.styles import ParagraphStyle                              # noqa: E402
from reportlab.lib.units import mm                                           # noqa: E402
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import asks as A            # noqa: E402
import calc as C            # noqa: E402
import common as K          # noqa: E402
import docs as DOC          # noqa: E402
import postprocess as PP    # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
GREEN = "#1f5c3a"
GREY = "#a7adb4"
RUNS = 24


def r10(x):
    return int(round(x / 10.0)) * 10


def r1(x):
    return round(x + 1e-9, 1)


def pct(a, n):
    return int(round(100.0 * a / n))


def abutting_share(B):
    """each store's open Christmas shifts that one of its own casuals could take by staying on: a casual whose
    rostered shift at the store ends as the open shift starts, or starts as it ends, and whom the call-in module's
    checks would offer it (availability, 11 hours a day, 38 a week, 10 hours between days), as counts"""
    fw = C.Forward(B, None)
    fw.reset()
    n = collections.Counter()
    a = collections.Counter()
    for ln in fw.open:
        n[ln["loc"]] += 1
        for pc in fw.order[ln["loc"]]:
            if fw.home[pc] != ln["loc"]:
                continue
            mine = fw.book.get((pc, ln["day"]), ())
            if any(l2 == ln["loc"] and (e2 == ln["s"] or s2 == ln["e"]) for l2, s2, e2 in mine) and fw.eligible(pc, ln):
                a[ln["loc"]] += 1
                break
    return {loc: (a[loc], n[loc]) for loc in B.locs}


def extension_split(rows):
    out = {}
    for name, sel in (("ext", lambda r: r["ext"]), ("other", lambda r: not r["ext"])):
        rs = [r for r in rows if sel(r)]
        out[name] = (len(rs), sum(r["y"] for r in rs))
    return out


def compute(target):
    B = C.load(target)
    G = C.golden(B, runs=RUNS)
    S = A.stops(B)
    ab = abutting_share(B)
    ex = extension_split(G["rows"])
    names = {loc: B.loc_store[loc] for loc in B.locs}
    order = sorted(B.locs, key=lambda l: -G["unworked"][l])
    n_open = collections.Counter(ln["loc"] for ln in B.xmas if ln["open"])
    figs = collections.OrderedDict()
    figs["call_store"] = names[G["store"]]
    figs["crew_hours"] = r10(G["crew_hours"])
    figs["crew_hours_exact"] = round(G["crew_hours"], 1)
    figs["crew_open_shifts"] = n_open[G["store"]]
    figs["offers_total"] = len(G["rows"])
    figs["ext_offers"] = ex["ext"][0]
    figs["ext_accepted"] = ex["ext"][1]
    figs["ext_rate_pct"] = pct(ex["ext"][1], ex["ext"][0])
    figs["other_offers"] = ex["other"][0]
    figs["other_accepted"] = ex["other"][1]
    figs["other_rate_pct"] = pct(ex["other"][1], ex["other"][0])
    for loc in B.locs:
        k = names[loc].lower().replace(" ", "_")
        figs["open_hours_%s" % k] = r10(G["open"][loc])
        figs["open_shifts_%s" % k] = n_open[loc]
        figs["abut_pct_%s" % k] = pct(ab[loc][0], ab[loc][1])
        figs["abut_shifts_%s" % k] = ab[loc][0]
        figs["unworked_%s" % k] = r10(G["unworked"][loc])
    figs["runner_up"] = names[order[1]]
    figs["runs"] = RUNS
    for loc in B.locs:
        k = names[loc].lower().replace(" ", "_")
        a = S["A_golden"][loc]
        figs["A_put_out_%s" % k] = r1(a["put_out"])
        figs["A_covered_%s" % k] = r1(a["covered"])
        figs["A_uncovered_%s" % k] = r1(a["uncovered"])
        b = S["B_golden"][loc]
        figs["B_imports_%s" % k] = r1(b["imports"])
        figs["B_exports_%s" % k] = r1(b["exports"])
    return B, G, S, ab, ex, figs, order, names, n_open


# ------------------------------------------------------------------------------------------------ the chart
def chart_png(G, order, names, crew_store):
    fig, ax = plt.subplots(figsize=(6.8, 2.75), dpi=220)
    vals = [G["unworked"][l] for l in order]
    cols = [GREEN if l == crew_store else "#b9bec4" for l in order]
    ys = list(range(len(order)))[::-1]
    ax.barh(ys, vals, color=cols, height=0.58, edgecolor="white", linewidth=1.0)
    for y, v, l in zip(ys, vals, order):
        ax.text(v + max(vals) * 0.012, y, "{:,}".format(r10(v)), va="center", fontsize=7.6,
                color="#1f2329", fontweight="bold" if l == crew_store else "normal")
    ax.set_yticks(ys)
    ax.set_yticklabels([names[l] for l in order], fontsize=8.2, color="#1f2329")
    ax.set_xlabel("Paid hours of open shifts left unfilled by call-in, 27 November to 24 December 2026 (forecast, "
                  "no crew placed)", fontsize=7.2, color="#4a5058")
    ax.set_xlim(0, max(vals) * 1.14)
    ax.xaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda x, _: "{:,.0f}".format(x)))
    ax.tick_params(axis="x", labelsize=7, colors="#4a5058")
    ax.tick_params(axis="y", length=0)
    ax.grid(axis="x", color="#e3e6ea", linewidth=0.6)
    ax.set_axisbelow(True)
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.spines["bottom"].set_color("#c9cdd2")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    buf.seek(0)
    return buf


def write_pdf(path, B, G, ab, ex, figs, order, names, n_open):
    st = DOC.styles("Carlito", "Carlito-Bold", size=9.3)
    small = ParagraphStyle("sm2", parent=st["small"], fontSize=7.3, leading=9.1)
    meta = ParagraphStyle("meta", parent=st["body"], fontSize=8.4, leading=10.6, spaceAfter=0)
    ctitle = ParagraphStyle("ct", parent=st["h2"], fontSize=9.6, spaceBefore=4, spaceAfter=0)
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=17 * mm, rightMargin=17 * mm, topMargin=12 * mm,
                            bottomMargin=13 * mm, title="Relief crew placement, Christmas 2026",
                            author="Workforce Planning", subject="", creator="", producer="")
    crew = G["store"]
    head = Table([[Paragraph("<b>To</b>  Belinda Friedman, General Manager, Retail Operations", meta),
                   Paragraph("<b>Date</b>  13 November 2026", meta)],
                  [Paragraph("<b>From</b>  Workforce Planning", meta),
                   Paragraph("<b>Ref</b>  WP-04 relief crew placement, Christmas 2026", meta)]],
                 colWidths=[100 * mm, 76 * mm])
    head.setStyle(TableStyle([("LINEBELOW", (0, 1), (-1, 1), 0.6, colors.HexColor("#a9b0b8")),
                              ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6), ("TOPPADDING", (0, 0), (-1, -1), 1.0),
                              ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    story = [head, Spacer(1, 5),
             Paragraph("The relief crew should go to %s this Christmas" % names[crew], st["h1"]),
             Paragraph("Place the 25-person relief crew at <b>%s</b> for the four trading weeks from Friday 27 November "
                       "to Thursday 24 December. It will work <b>%s hours</b> there, the paid hours of the %d shifts "
                       "%s's Christmas roster lists open. Of all eight stores, %s's open shifts are the ones our "
                       "call-ins would most leave unworked, well clear of any other store's."
                       % (names[crew], "{:,}".format(figs["crew_hours"]), n_open[crew], names[crew], names[crew]),
                       st["body"]),
             Paragraph("Open shifts call-ins would leave unworked if the crew went nowhere", ctitle),
             Image(chart_png(G, order, names, crew), width=176 * mm, height=71 * mm),
             Paragraph("<b>Why not Werribee.</b> Werribee lists by far the most open hours, %s, and on last year's "
                       "fill rates it looks the obvious place for the crew again. But its Christmas roster puts the "
                       "casuals it took on in February mostly on 8-to-2 and 2-to-6 shifts and lists the extra "
                       "shifts open from 2 pm and from 6 pm until midnight, straight after them: %d%% of its open "
                       "shifts could be taken by one of its own casuals staying on from a rostered shift. Box Hill "
                       "(%d%%) and Altona (%d%%) have done the same with their evening uplift. Casuals take that kind "
                       "of offer far more readily than one that needs a separate trip in: since August 2025, %d%% of "
                       "the %s offers that carried on from the casual's own shift at the store were accepted, against "
                       "%d%% of the other %s. Our call-ins should fill most of Werribee's Christmas uplift without "
                       "the crew."
                       % ("{:,}".format(figs["open_hours_werribee"]), figs["abut_pct_werribee"],
                          figs["abut_pct_box_hill"], figs["abut_pct_altona"], figs["ext_rate_pct"],
                          "{:,}".format(figs["ext_offers"]), figs["other_rate_pct"], "{:,}".format(figs["other_offers"])),
                       st["body"]),
             Paragraph("<b>Why %s.</b> Its uplift is daytime work, 9 am to 3 pm on weekdays and morning shifts at "
                       "weekends. Plenty of casuals are free for it, most of them Box Hill's and Ringwood's university "
                       "students on their summer break, but %s of its open shifts could be taken by a casual staying "
                       "on from a rostered shift, so every one needs someone to make a separate trip in, most of them "
                       "to another store, in the busiest month of the year, and few casuals do."
                       % (names[crew], "none" if figs["abut_pct_camberwell"] == 0 else
                          "only %d%%" % figs["abut_pct_camberwell"]), st["body"])]
    rows = [["Store", "Open shifts", "Open paid hours", "Could be taken by staying on", "Forecast unworked hours"]]
    for l in order:
        k = names[l].lower().replace(" ", "_")
        rows.append([names[l], str(n_open[l]), "{:,}".format(figs["open_hours_%s" % k]), "%d%%" % figs["abut_pct_%s" % k],
                     "{:,}".format(figs["unworked_%s" % k])])
    t = Table([[Paragraph(c, st["cellb"] if i == 0 else st["cell"]) for c in r] for i, r in enumerate(rows)],
              colWidths=[30 * mm, 22 * mm, 30 * mm, 50 * mm, 44 * mm])
    t.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#a9b0b8")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e5ebe7")),
                           ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#eef5f0")),
                           ("TOPPADDING", (0, 0), (-1, -1), 1.0), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.0)]))
    story += [Spacer(1, 3), t,
              Paragraph("Source: KWS extract WR-26-1187 (12 November 2026); Christmas rosters as published 12 November "
                        "2026, 10:00. Hours are rounded to the nearest 10.", small),
              Spacer(1, 3),
              Paragraph("Forecast: each casual's answer to a call-in offer fitted as a yes-or-no choice on all %s offers "
                        "in the extract (the shift's paid hours, time and notice, a sister-store offer, hours already "
                        "rostered that week, the Christmas weeks, the casual's age band and home store, and whether "
                        "the offer carried on from the casual's own shift), then run %d times through the call-in "
                        "module's offer order and limits over the published rosters, every store's open shifts posted "
                        "at 09:00 on 16 November and the crew at no store, as WP-04 section 4 counts it. Paid hours "
                        "are net of the unpaid 30-minute meal break on shifts over five hours (EA 2024). Last "
                        "Christmas's call-in tables are in relief_crew_xmas2026.xlsx."
                        % ("{:,}".format(figs["offers_total"]), RUNS), small)]

    def foot(canvas, d):
        canvas.saveState()
        canvas.setFont("Carlito", 7.2)
        canvas.setFillColor(colors.HexColor("#5f6670"))
        canvas.drawString(17 * mm, 7 * mm, "Kellaway's Workforce Planning")
        canvas.drawRightString(A4[0] - 17 * mm, 7 * mm, "Page 1 of 1")
        canvas.restoreState()

    doc.build(story, onFirstPage=foot, onLaterPages=foot)
    PP.normalize_pdf(path, b"D:20261113161900+11'00'", b"D:20261113161900+11'00'")


# ------------------------------------------------------------------------------------------------ the workbook
def _head(ws, r, cols):
    for j, h in enumerate(cols, 1):
        c = ws.cell(row=r, column=j, value=h)
        c.font = Font(name="Calibri", bold=True)
        c.fill = PatternFill("solid", fgColor="E5EBE7")
        c.alignment = Alignment(wrap_text=True, vertical="top")


def write_xlsx(path, B, G, S, ab, ex, figs, order, names, n_open):
    import numpy as np
    wb = Workbook()
    bold = Font(name="Calibri", bold=True)
    ws = wb.active
    ws.title = "Placement"
    ws["A1"] = "Relief crew placement, Christmas 2026: the crew to %s" % names[G["store"]]
    ws["A1"].font = Font(name="Calibri", size=13, bold=True)
    ws["A2"] = "Crew placed at"
    ws["B2"] = names[G["store"]]
    ws["B2"].font = bold
    ws["A3"] = "Crew hours, paid hours of the store's open shifts (nearest 10)"
    ws["B3"] = figs["crew_hours"]
    ws["B3"].number_format = "#,##0"
    ws["B3"].font = bold
    cols = ["Store", "Open shifts", "Open paid hours", "Open shifts a casual could take by staying on",
            "Share", "Forecast unworked hours (mean of %d runs)" % RUNS, "Rank"]
    _head(ws, 5, cols)
    for i, l in enumerate(order):
        a, n = ab[l]
        vals = [names[l], n_open[l], G["open"][l], a, a / n, G["unworked"][l], i + 1]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=6 + i, column=j, value=round(v, 4) if isinstance(v, float) else v)
            c.number_format = {3: "#,##0.0", 5: "0%", 6: "#,##0"}.get(j, "General")
    r = 6 + len(order)
    ws.cell(row=r, column=1, value="All stores").font = bold
    ws.cell(row=r, column=2, value=sum(n_open[l] for l in order)).font = bold
    c = ws.cell(row=r, column=3, value=round(sum(G["open"][l] for l in order), 1))
    c.number_format, c.font = "#,##0.0", bold
    ws.cell(row=r, column=4, value=sum(ab[l][0] for l in order)).font = bold
    c = ws.cell(row=r, column=6, value=round(sum(G["unworked"][l] for l in order), 1))
    c.number_format, c.font = "#,##0", bold
    ws.cell(row=r + 2, column=1, value="Open paid hours to one decimal; forecast hours to the nearest hour here, the "
                                       "brief rounds both to the nearest 10.")
    ws.freeze_panes = "B6"
    ws2 = wb.create_sheet("Take-up model")
    ws2["A1"] = "Call-in offers 29/08/2025 to 12/11/2026, by whether the offer carried on from the casual's own " \
                "rostered or booked shift at that store"
    ws2["A1"].font = bold
    _head(ws2, 3, ["Offer", "Offers", "Accepted", "Accepted share"])
    for i, (name, key) in enumerate((("Carries on from own shift", "ext"), ("Needs a separate trip", "other"))):
        n, acc = ex[key]
        for j, v in enumerate([name, n, acc, acc / n], 1):
            c = ws2.cell(row=4 + i, column=j, value=round(v, 4) if isinstance(v, float) else v)
            c.number_format = {2: "#,##0", 3: "#,##0", 4: "0.0%"}.get(j, "General")
    names_b = (["Constant", "Paid hours of the shift", "Sister-store offer", "Hours rostered that week, over 10",
                "Ends after 21:30", "Starts before 07:00", "Sunday", "Under four hours' notice", "Offered the shift before",
                "Christmas weeks", "Aged under 18", "Aged 18 to 24"] +
               ["Home store %s" % B.loc_store[l] for l in B.locs[1:]] + ["Carries on from own shift",
                                                                        "Another shift that day, not carrying on"])
    ws2["A7"] = "Logit of acceptance, fitted on every offer (base: aged 25 and over, home store %s)" % B.loc_store[B.locs[0]]
    ws2["A7"].font = bold
    _head(ws2, 8, ["Term", "Coefficient", "Standard error"])
    se = np.sqrt(np.diag(G["model"]["cov"]))
    for i, (nm, b, s) in enumerate(zip(names_b, G["model"]["b"], se)):
        ws2.cell(row=9 + i, column=1, value=nm)
        ws2.cell(row=9 + i, column=2, value=round(float(b), 3)).number_format = "0.000"
        ws2.cell(row=9 + i, column=3, value=round(float(s), 3)).number_format = "0.000"
    ws3 = wb.create_sheet("Last Christmas call-ins")
    ws3["A1"] = "Shifts put out to call-in for 28/11/2025 to 24/12/2025, paid hours"
    ws3["A1"].font = bold
    _head(ws3, 3, ["Store", "Hours put out to call-in", "Covered", "Left uncovered"])
    tot = [0.0, 0.0, 0.0]
    srt = sorted(B.locs, key=lambda l: names[l])
    for i, l in enumerate(srt):
        k = names[l].lower().replace(" ", "_")
        vals = [figs["A_put_out_%s" % k], figs["A_covered_%s" % k], figs["A_uncovered_%s" % k]]
        ws3.cell(row=4 + i, column=1, value=names[l])
        for j, v in enumerate(vals):
            ws3.cell(row=4 + i, column=2 + j, value=v).number_format = "#,##0.0"
            tot[j] += v
    ws3.cell(row=4 + len(srt), column=1, value="All stores").font = bold
    for j, v in enumerate(tot):
        c = ws3.cell(row=4 + len(srt), column=2 + j, value=round(v, 1))
        c.number_format, c.font = "#,##0.0", bold
    ws3.cell(row=6 + len(srt), column=1, value="Each shift counts once at its final outcome: a shift raised again after "
                                               "its list ran out, raised again after the casual who took it gave it "
                                               "back, or vacated by an absence reported twice, is one shift, and so is "
                                               "an open line the 2025 roster import raised twice (open lines counted "
                                               "from the 2025 roster export); the identical open lines a roster lists "
                                               "are each a shift. A shift is covered by the casual who works it. "
                                               "Werribee's Christmas uplift went to last year's relief crew and was "
                                               "never put out.")
    ws3.freeze_panes = "B4"
    ws4 = wb.create_sheet("Cover across stores")
    ws4["A1"] = "Call-in hours covered across stores, 28/11/2025 to 24/12/2025, paid hours"
    ws4["A1"].font = bold
    _head(ws4, 3, ["Store", "Its call-ins covered by other stores' casuals", "Its own casuals' cover at other stores"])
    tb = [0.0, 0.0]
    for i, l in enumerate(srt):
        k = names[l].lower().replace(" ", "_")
        ws4.cell(row=4 + i, column=1, value=names[l])
        for j, v in enumerate([figs["B_imports_%s" % k], figs["B_exports_%s" % k]]):
            ws4.cell(row=4 + i, column=2 + j, value=v).number_format = "#,##0.0"
            tb[j] += v
    ws4.cell(row=4 + len(srt), column=1, value="All stores").font = bold
    for j, v in enumerate(tb):
        c = ws4.cell(row=4 + len(srt), column=2 + j, value=round(v, 1))
        c.number_format, c.font = "#,##0.0", bold
    ws4.cell(row=6 + len(srt), column=1, value="A casual's own store is their home store on the shift date, from the "
                                               "payroll cost-centre history, each cost centre counted from the start "
                                               "of its engagement: a move between stores before April 2026 was "
                                               "recorded as a re-engagement under a new employee number, its cost "
                                               "centre dated from the day the move was keyed. Ballarat offers to no "
                                               "other store's casuals.")
    ws4.freeze_panes = "B4"
    ws5 = wb.create_sheet("Notes")
    notes = [
        "Placement standard WP-04 v4: the crew goes to the store where it will work the most hours that would otherwise "
        "go unworked; it works every shift the store's published Christmas roster lists open (sections 3 to 5).",
        "Open shifts: lines with no employee in the Christmas rosters published 12/11/2026 10:00.",
        "Paid hours: shift length less the unpaid 30-minute meal break on shifts over five hours (EA 2024).",
        "Forecast: the fitted choice run through the KWS offer order, eligibility checks and offer hours, every store's "
        "open shifts posted 16/11/2026 09:00 with the crew at no store, %d runs; the mean of the runs." % RUNS,
        "Sources: KWS extract WR-26-1187 (12/11/2026); payroll casual register (12/11/2026).",
        "Prepared by Workforce Planning, 13/11/2026.",
    ]
    ws5["A1"] = "Notes"
    ws5["A1"].font = bold
    for i, x in enumerate(notes):
        ws5.cell(row=3 + i, column=1, value=x)
    widths = {"Placement": [34, 12, 16, 22, 14, 22, 8], "Take-up model": [40, 14, 14, 14],
              "Last Christmas call-ins": [18, 22, 14, 16], "Cover across stores": [18, 30, 30], "Notes": [120]}
    for w in wb.worksheets:
        for j, wd in enumerate(widths[w.title], 1):
            w.column_dimensions[get_column_letter(j)].width = wd
    wb.save(path)
    PP.normalize_ooxml(path, "Beth Hill", "Beth Hill", "2026-11-13T00:12:00Z", "2026-11-13T05:18:00Z",
                       title="Relief crew placement, Christmas 2026", excel=True, company="Kellaway's")


def write_goldens(task, target):
    B, G, S, ab, ex, figs, order, names, n_open = compute(target)
    gd = os.path.join(task, "golden")
    os.makedirs(gd, exist_ok=True)
    write_pdf(os.path.join(gd, K.DELIVERABLES[0]), B, G, ab, ex, figs, order, names, n_open)
    write_xlsx(os.path.join(gd, K.DELIVERABLES[1]), B, G, S, ab, ex, figs, order, names, n_open)
    with open(os.path.join(HERE, "figures.json"), "w", encoding="utf-8") as fh:
        json.dump(figs, fh, indent=1)
        fh.write("\n")
    stamp = dt.datetime(2026, 11, 13, 16, 20).timestamp()
    for f in os.listdir(gd):
        os.utime(os.path.join(gd, f), (stamp, stamp))
    return dict(B=B, G=G, S=S, ab=ab, ex=ex, figs=figs, order=order, names=names, n_open=n_open)
