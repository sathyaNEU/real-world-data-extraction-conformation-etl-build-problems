"""Independent verifier for task155: reads only the shipped bundle (pyarrow for the CSVs, openpyxl for the Christmas
workbook), shares no code with calc.py or asks.py, and recomputes every figure figures.json holds that the bundle
fixes exactly: the crew hours and open shifts, the extension split of the offer history, each store's open shifts
against a casual's own shift, and both ask tables.

    python3 verify.py <target> <figures.json>"""
from __future__ import annotations

import collections
import datetime as dt
import json
import math
import os
import sys
from zoneinfo import ZoneInfo

import numpy as np
import pyarrow as pa
import pyarrow.csv as pv
from openpyxl import load_workbook
from scipy.optimize import minimize

MEL = ZoneInfo("Australia/Melbourne")
EPOCH = dt.date(2025, 8, 29)


def tab(target, name):
    path = os.path.join(target, name)
    with open(path, encoding="utf-8") as fh:
        cols = fh.readline().strip().split(",")
    t = pv.read_csv(path, convert_options=pv.ConvertOptions(column_types={c: pa.string() for c in cols},
                                                            strings_can_be_null=False))
    return {c: [("" if v is None else str(v)) for v in t.column(c).to_pylist()] for c in t.column_names}


def local(ts):
    t = dt.datetime.fromisoformat(ts.replace("Z", "+00:00")).astimezone(MEL)
    return t.replace(tzinfo=None)


def mins(t):
    return (t.date() - EPOCH).days * 1440 + t.hour * 60 + t.minute + t.second / 60


def paid(span_min):
    return span_min - 30 if span_min > 300 else span_min


def clock(s):
    s = s.strip()
    if s.endswith(("AM", "PM")):
        hh, mm = s[:-3].split(":")
        return (int(hh) % 12 + (12 if s.endswith("PM") else 0)) * 60 + int(mm)
    hh, mm = s.split(":")[:2]
    return int(hh) * 60 + int(mm)


def main(target, figfile):
    F = json.load(open(figfile))
    bad = []

    def same(label, got, want, tol=0.05):
        ok = abs(float(got) - float(want)) <= tol
        print("%-4s %-58s %s / %s" % ("ok" if ok else "FAIL", label, got, want))
        if not ok:
            bad.append(label)

    stores = tab(target, "stores.csv")
    name = dict(zip(stores["kws_location"], stores["store_name"]))
    loc_by_name = {v.upper(): k for k, v in name.items()}
    cc_loc = dict(zip(stores["cost_centre"], stores["kws_location"]))
    cluster = dict(zip(stores["kws_location"], stores["callin_cluster"]))

    # the Christmas rosters
    wb = load_workbook(os.path.join(target, "christmas_roster_2026_published.xlsx"), read_only=True, data_only=True)
    xmas = []
    for ws in wb.worksheets:
        loc = loc_by_name[ws.title.upper()]
        for row in ws.iter_rows(min_row=7, values_only=True):
            if row[0] is None:
                continue
            day = (row[0].date() - EPOCH).days
            xmas.append((loc, day, clock(row[3]), clock(row[4]), row[6], row[7], row[8]))
    open_h = collections.Counter()
    open_n = collections.Counter()
    for loc, day, s, e, pn, typ, st in xmas:
        if st == "OPEN":
            open_h[loc] += paid(e - s) / 60
            open_n[loc] += 1
    for loc in name:
        k = name[loc].lower().replace(" ", "_")
        same("open hours %s (nearest 10)" % name[loc], int(round(open_h[loc] / 10.0)) * 10, F["open_hours_%s" % k], 0)
        same("open shifts %s" % name[loc], open_n[loc], F["open_shifts_%s" % k], 0)
    crew_loc = loc_by_name[F["call_store"].upper()]
    same("crew hours (nearest 10)", int(round(open_h[crew_loc] / 10.0)) * 10, F["crew_hours"], 0)
    # identities
    ppl = tab(target, "kws_people_20261112.csv")
    pay = tab(target, "payroll_casual_register_20261112.csv")
    prev = {int(n): int(p) for n, p in zip(pay["employee_no"], pay["previous_employee_no"]) if p}
    periods = collections.defaultdict(list)
    for n, f, t, cc in zip(pay["employee_no"], pay["cc_effective_from"], pay["cc_effective_to"], pay["cost_centre"]):
        periods[int(n)].append((dt.date.fromisoformat(f), dt.date.fromisoformat(t) if t else dt.date(2100, 1, 1), cc_loc[cc]))
    chain = {}
    for pc, n in zip(ppl["person_code"], ppl["payroll_no"]):
        c = [int(n)]
        while c[0] in prev:
            c = [prev[c[0]]] + c
        chain[pc] = c

    def number_on(pc, d):
        for n in chain[pc]:
            if any(f <= d <= t for f, t, _ in periods[n]):
                return n
        return chain[pc][-1]

    def home_on(pc, d):
        for n in chain[pc]:
            for f, t, loc in periods[n]:
                if f <= d <= t:
                    return loc
        return None

    # open Christmas shifts one of the store's own casuals could take by staying on, the module's checks passed
    prof = {}
    for p in json.load(open(os.path.join(target, "kws_availability_20261112.json"), encoding="utf-8"))["people"]:
        if p["profiles"] and p["profiles"][-1]["effective_to"] is None:
            prof[p["person_code"]] = {k: [(clock(a), clock(b)) for a, b in w] for k, w in p["profiles"][-1]["weekly"].items()}
    active = {pc: (int(n), home) for pc, n, home, st in zip(ppl["person_code"], ppl["payroll_no"], ppl["home_location"],
                                                            ppl["status"]) if st == "ACTIVE"}
    by_no = {n: pc for pc, (n, home) in active.items()}
    days = collections.defaultdict(list)
    for loc, day, s, e, pn, typ, st in xmas:
        if st != "OPEN" and pn is not None and int(pn) in by_no:
            days[(by_no[int(pn)], day)].append((loc, s, e))
    wd = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
    for loc in name:
        k = name[loc].lower().replace(" ", "_")
        a = 0
        for l2, day, s, e, pn, typ, st in xmas:
            if l2 != loc or st != "OPEN":
                continue
            d = EPOCH + dt.timedelta(days=day)
            for pc, (no, home) in active.items():
                if home != loc:
                    continue
                mine = days.get((pc, day), [])
                if not any(x[0] == loc and (x[2] == s or x[1] == e) for x in mine):
                    continue
                if not any(lo <= s and e <= hi for lo, hi in prof.get(pc, {}).get(wd[d.weekday()], [])):
                    continue
                if any(x[1] < e and s < x[2] for x in mine):
                    continue
                if sum(paid(x[2] - x[1]) for x in mine) + paid(e - s) > 660:
                    continue
                wk0 = (day // 7) * 7
                if sum(paid(x[2] - x[1]) for dd in range(wk0, wk0 + 7) for x in days.get((pc, dd), [])) + paid(e - s) > 38 * 60:
                    continue
                if any(day * 1440 + s - ((day - 1) * 1440 + x[2]) < 600 for x in days.get((pc, day - 1), [])):
                    continue
                if any((day + 1) * 1440 + x[1] - (day * 1440 + e) < 600 for x in days.get((pc, day + 1), [])):
                    continue
                a += 1
                break
        same("open shifts a casual could take by staying on, %s" % name[loc], a, F["abut_shifts_%s" % k], 0)

    # the register and the rosters
    reg = tab(target, "callin_register_20250829_20261112.csv")
    R = []
    for i in range(len(reg["callin_id"])):
        a, b = local(reg["shift_start_utc"][i]), local(reg["shift_end_utc"][i])
        R.append(dict(id=reg["callin_id"][i], loc=reg["location_id"][i], d=a.date(), s=a.hour * 60 + a.minute,
                      e=(b.hour * 60 + b.minute) or 1440, span=(b - a).total_seconds() / 60, reason=reg["reason"][i],
                      absent=reg["absent_payroll_no"][i], raised=local(reg["raised_at"][i]),
                      closed=local(reg["closed_at"][i]), outcome=reg["outcome"][i], by=reg["filled_by"][i]))
    vacated = {(x["loc"], x["d"], x["s"], x["e"], int(x["absent"])) for x in R if x["absent"]}
    lines = collections.defaultdict(list)
    leg = tab(target, "roster_export_legacy_20250829_20260226.csv")
    for i in range(len(leg["Store"])):
        if leg["Status"][i] or not leg["Emp No"][i]:
            continue
        d = dt.datetime.strptime(leg["Date"][i], "%d/%m/%Y").date()
        key = (loc_by_name[leg["Store"][i]], d, clock(leg["Start"][i]), clock(leg["Finish"][i]), int(leg["Emp No"][i]))
        if key not in vacated:
            lines[(key[4], d)].append(key[:4])
    kws = tab(target, "roster_lines_kws_20260227_20261112.csv")
    for i in range(len(kws["location_id"])):
        if kws["line_status"][i] != "ROSTERED":
            continue
        d = dt.date.fromisoformat(kws["shift_date"][i])
        key = (kws["location_id"][i], d, clock(kws["start_time"][i]), clock(kws["end_time"][i]), int(kws["payroll_no"][i]))
        if key not in vacated:
            lines[(key[4], d)].append(key[:4])
    person_of = {n: pc for pc, c in chain.items() for n in c}
    drop_at = {}
    for x in sorted(R, key=lambda x: x["raised"]):
        if x["reason"] == "ABS" and x["absent"] and int(x["absent"]) in person_of:
            drop_at.setdefault((person_of[int(x["absent"])], x["loc"], x["d"], x["s"], x["e"]), x["raised"])
    fills = collections.defaultdict(list)
    for x in R:
        if x["outcome"] == "FILLED":
            td = drop_at.get((x["by"], x["loc"], x["d"], x["s"], x["e"]))
            fills[(x["by"], x["d"])].append((x["loc"], x["s"], x["e"], x["closed"], td if td and td > x["closed"] else None))

    # the extension split
    off = tab(target, "callin_offers_20250829_20261112.csv")
    n_ext = a_ext = n_oth = a_oth = 0
    for i in range(len(off["offer_id"])):
        pc = off["person_code"][i]
        d = dt.date.fromisoformat(off["shift_date"][i])
        s, e = clock(off["shift_start"][i]), clock(off["shift_end"][i])
        loc = off["location_id"][i]
        t = local(off["offered_at"][i])
        mine = list(lines.get((number_on(pc, d), d), []))
        for (l2, s2, e2, ta, td) in fills.get((pc, d), []):
            if ta <= t and (td is None or td > t) and not (l2 == loc and s2 == s and e2 == e):
                mine.append((l2, d, s2, e2))
        ext = any(m[0] == loc and (m[-1] == s or m[-2] == e) for m in mine)
        acc = off["response"][i] == "ACCEPTED"
        if ext:
            n_ext += 1
            a_ext += acc
        else:
            n_oth += 1
            a_oth += acc
    same("offers carrying on from the casual's own shift", n_ext, F["ext_offers"], 0)
    same("of them accepted", a_ext, F["ext_accepted"], 0)
    same("other offers", n_oth, F["other_offers"], 0)
    same("of them accepted", a_oth, F["other_accepted"], 0)

    # ask A and ask B, on a path of their own: each line's call-ins walked in the order KWS raised them, a re-raise
    # at 07:00 taking the next line of its shift still waiting, a second report of an absence joining its line, a
    # give-back's ABS call-in carrying on the line its casual had filled
    win = (dt.date(2025, 11, 28), dt.date(2025, 12, 24))
    W_ = sorted([x for x in R if win[0] <= x["d"] <= win[1]], key=lambda x: (x["raised"], int(x["id"])))
    walked = []
    waiting = collections.defaultdict(list)       # (shift, reason, absent) -> lines whose last call-in went unfilled
    reported = {}                                 # (shift, reason, absent) -> the line an absence or leave vacated
    taken = {}                                    # (shift, payroll no) -> the line that casual filled
    for x in W_:
        key = (x["loc"], x["d"], x["s"], x["e"])
        tag = (key, x["reason"], x["absent"])
        at7 = x["raised"].hour == 7 and x["raised"].minute == 0 and x["raised"].second == 0
        line = None
        if x["reason"] == "ABS" and (key, x["absent"]) in taken:
            line = taken.pop((key, x["absent"]))          # given back: the first fill no longer covers it
            line["fill"] = None
        elif at7 and waiting[tag]:
            line = waiting[tag].pop(0)
        elif x["reason"] != "OPN" and tag in reported:
            line = reported[tag]                           # the same absence, reported again
        if line is None:
            line = dict(loc=x["loc"], d=x["d"], h=paid(x["span"]) / 60, fill=None)
            walked.append(line)
            if x["reason"] != "OPN":
                reported[tag] = line
        if x["outcome"] == "FILLED":
            line["fill"] = x["by"]
            taken[(key, str(number_on(x["by"], x["d"])))] = line
        elif line["fill"] is None and all(w is not line for w in waiting[tag]):
            waiting[tag].append(line)
    put = collections.Counter()
    cov = collections.Counter()
    imp = collections.Counter()
    exp = collections.Counter()
    for line in walked:
        put[line["loc"]] += line["h"]
        if line["fill"]:
            cov[line["loc"]] += line["h"]
            home = home_on(line["fill"], line["d"])
            if home != line["loc"]:
                imp[line["loc"]] += line["h"]
                exp[home] += line["h"]
    for loc in name:
        k = name[loc].lower().replace(" ", "_")
        same("ask A put out %s" % name[loc], round(put[loc], 1), F["A_put_out_%s" % k])
        same("ask A covered %s" % name[loc], round(cov[loc], 1), F["A_covered_%s" % k])
        same("ask A uncovered %s" % name[loc], round(put[loc] - cov[loc], 1), F["A_uncovered_%s" % k])
        same("ask B from other stores %s" % name[loc], round(imp[loc], 1), F["B_imports_%s" % k])
        same("ask B at other stores %s" % name[loc], round(exp[loc], 1), F["B_exports_%s" % k])

    print("verify: %d failed" % len(bad))
    return 0 if not bad else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1], sys.argv[2]))
