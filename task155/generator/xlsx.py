"""The pack's two workbooks: the published Christmas 2026 rosters (one sheet per store, as the workforce system's
roster report saves them) and Store Operations' Christmas trading plan."""
from __future__ import annotations

import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import common as K
import world as W

THIN = Side(style="thin", color="B7BDC5")
HEAD_FILL = PatternFill("solid", fgColor="DCE3EA")
OPEN_FILL = PatternFill("solid", fgColor="FCEFD9")


def write_christmas(path, Wd):
    import records as R
    wb = Workbook()
    wb.remove(wb.active)
    n_lines = 0
    for s in sorted(W.STORES, key=lambda s: K.STORE_NO[s]):
        ws = wb.create_sheet(W.STORE_NAMES[s])
        ws["A1"] = "Kellaway's"
        ws["A1"].font = Font(name="Calibri", size=11, bold=True)
        ws["A2"] = "Christmas roster 2026: %s (%s)" % (W.STORE_NAMES[s], K.LOC[s])
        ws["A2"].font = Font(name="Calibri", size=14, bold=True)
        ws["A3"] = "Roster weeks Friday 27/11/2026 to Thursday 24/12/2026"
        ws["A4"] = "Published 12/11/2026 10:00"
        head = ["Date", "Day", "Shift", "Start", "Finish", "Paid hrs", "Emp No", "Type", "Status"]
        for j, h in enumerate(head, 1):
            c = ws.cell(row=6, column=j, value=h)
            c.font = Font(name="Calibri", size=11, bold=True)
            c.fill = HEAD_FILL
            c.border = Border(bottom=THIN)
            c.alignment = Alignment(horizontal="center")
        r = 7
        for ln in sorted(Wd.roster_p[s], key=lambda l: (l["day"], l["start"], l["end"], l["code"], l["who"] is None)):
            d = W.ddate(ln["day"])
            who = ln["who"]
            vals = [d, d.strftime("%a"), ln["code"], K.hhmm(ln["start"]), K.hhmm(ln["end"]),
                    round(K.paid_hours(ln["start"], ln["end"]), 2),
                    R.payroll_no(Wd, who, d) if who is not None else None, R.emp_type(Wd, who),
                    "OPEN" if who is None else "ROSTERED"]
            for j, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=j, value=v)
                if j == 1:
                    c.number_format = "dd/mm/yyyy"
                if j == 6:
                    c.number_format = "0.00"
                if who is None:
                    c.fill = OPEN_FILL
            r += 1
            n_lines += 1
        for j, w in enumerate([12, 6, 7, 8, 8, 9, 10, 6, 11], 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = "A7"
    wb.save(path)
    return n_lines


PLAN_SALES = {"W1": 1182, "W2": 1394, "W3": 1268, "W4": 1041, "E1": 1307, "E2": 1226, "E3": 1119, "RG": 1076}
WEEK_SHAPE = [0.93, 0.99, 1.04, 1.21]


def write_plan(path):
    """Store Operations' Christmas trading plan: sales budget and customer numbers by store and roster week"""
    rng = W.stream(21)
    wb = Workbook()
    ws = wb.active
    ws.title = "Christmas 2026"
    ws["A1"] = "Kellaway's Store Operations"
    ws["A1"].font = Font(name="Calibri", size=11, bold=True)
    ws["A2"] = "Christmas trading plan 2026: sales budget ($000) and customer count, by roster week"
    ws["A2"].font = Font(name="Calibri", size=13, bold=True)
    ws["A3"] = "Version 2, 30/10/2026. Budget set against 2025 actuals; customer counts from the front-end systems."
    weeks = [W.PEAK26[0] + dt.timedelta(days=7 * k) for k in range(4)]
    head = ["Store", "Cluster"] + ["Sales w/c %s" % w.strftime("%d/%m") for w in weeks] + \
        ["Customers w/c %s" % w.strftime("%d/%m") for w in weeks] + ["Trading hours note"]
    for j, h in enumerate(head, 1):
        c = ws.cell(row=5, column=j, value=h)
        c.font = Font(name="Calibri", size=11, bold=True)
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    notes = {"W2": "to midnight Mon to Sun from 27/11", "E2": "online picking 10:00 to 16:00 weekdays",
             "RG": "", "W1": "", "W3": "", "W4": "", "E1": "", "E3": ""}
    r = 6
    for s in sorted(W.STORES, key=lambda s: K.STORE_NO[s]):
        base = PLAN_SALES[s]
        sales = [round(base * f * rng.uniform(0.985, 1.015), 1) for f in WEEK_SHAPE]
        basket = rng.uniform(38.5, 47.0)
        cust = [int(round(v * 1000 / basket)) for v in sales]
        vals = [W.STORE_NAMES[s], K.CLUSTER_NAME[W.STORES[s]["cluster"]] or "Regional"] + sales + cust + [notes[s]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            if 3 <= j <= 6:
                c.number_format = "#,##0.0"
            if 7 <= j <= 10:
                c.number_format = "#,##0"
        r += 1
    for j, w in enumerate([14, 9, 11, 11, 11, 11, 12, 12, 12, 12, 34], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    wb.save(path)
    return 8
