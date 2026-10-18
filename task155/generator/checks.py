"""The generator's assertions on the built pack: the three rules executed on the clean truth, the ladder, every
semantic fork, the closure axes, the asks and their devices, the separation line, the input gates and the realism
sweep. Every figure is recomputed from the shipped bundle; the world's own truth (the Christmas cascade with each
casual's real answers) is run only to show the golden forecast picks the store the world would."""
from __future__ import annotations

import collections
import copy
import datetime as dt
import json
import math
import os
import re
import statistics
import sys
from multiprocessing import Pool

import numpy as np

import asks as A
import calc as C
import common as K
import docs as DOC
import world as W

RESULTS = []


def check(label, ok, detail=""):
    RESULTS.append((label, bool(ok), detail))
    print("%-4s %-52s %s" % ("ok" if ok else "FAIL", label, detail))


def ratio(u, a, b):
    return u[a] / u[b] if u[b] > 0 else float("inf")


# ------------------------------------------------------------------------------------------------ the world's truth
def _xmas_run(r):
    import records as R
    Wd = _xmas_run.Wd
    nd, vac, absd = W.make_needs(W.stream(1000 + r), Wd.roster_p, Wd.cas, Wd.perms, set(R.XMAS_DAYS),
                                 open_post=W.tmin(W.POST26))
    sp = W.Cascade(Wd.cas, Wd.roster_p, vac, absd, W.stream(2000 + r), W.stream(3000 + r), W.stream(4000 + r))
    sp.run(nd, log=False)
    st = W.settle(nd)
    u = collections.Counter()
    for k, fin in st.items():
        root = nd[k]
        if root["kind"] == "OPEN" and fin["filled"] is None:
            u[K.LOC[root["store"]]] += W.paid(root["e"] - root["s"]) / 60.0
    return dict(u)


def world_truth(runs=32):
    import records as R
    Wd = R.prepare()
    _xmas_run.Wd = Wd
    with Pool(4) as pool:
        res = pool.map(_xmas_run, range(1, runs + 1))
    locs = [K.LOC[s] for s in W.STORES]
    mean = {l: float(np.mean([r.get(l, 0.0) for r in res])) for l in locs}
    med = {l: float(np.median([r.get(l, 0.0) for r in res])) for l in locs}
    return Wd, mean, med, res


def world_asks(Wd, B):
    """last Christmas's call-ins as the world ran them: every line put out once (an absence reported twice is one
    line), covered when the last fill on it held (a give-back's re-raise carries the line on), each filler's home
    store on the shift date"""
    final = W.settle(Wd.needs)
    lines = collections.OrderedDict()
    for nid, fin in final.items():
        nd = Wd.needs[nid]
        d = W.ddate(nd["day"])
        if not (A.WINDOW[0] <= d <= A.WINDOW[1]) or not nd["postings"]:
            continue
        key = (nd["line"], nd["kind"], nd["absent"])
        lines.setdefault(key, []).append((nd, fin))
    put = collections.Counter()
    cov = collections.Counter()
    imp = collections.Counter()
    exp = collections.Counter()
    for key, lst in lines.items():
        nd = lst[0][0]
        loc = K.LOC[nd["store"]]
        h = W.paid(nd["e"] - nd["s"]) / 60.0
        put[loc] += h
        filled = [fin["filled"] for _, fin in lst if fin["filled"] is not None]
        if filled:
            cov[loc] += h
            home = K.LOC[Wd.cas[filled[0][0]].home(W.ddate(nd["day"]))]
            if home != loc:
                imp[loc] += h
                exp[home] += h
    tA = {l: dict(put_out=put[l], covered=cov[l], uncovered=put[l] - cov[l]) for l in B.locs}
    tB = {l: dict(imports=imp[l], exports=exp[l]) for l in B.locs}
    return tA, tB


# ------------------------------------------------------------------------------------------------ routes on the bundle
def routes(B, rows):
    out = collections.OrderedDict()
    oh = C.open_hours(B)
    out["open hours"] = oh
    ordinary = C.register_shares(B, lambda d: not C.in_x25(d))
    out["R0 ordinary-week register fill"] = {l: oh[l] * (1 - ordinary[l]) for l in B.locs}
    last = C.register_shares(B, C.in_x25)
    out["R1 last Christmas register fill"] = {l: oh[l] * (1 - last[l]) for l in B.locs}
    ordw = [(dt.date(2025, 8, 29), dt.date(2025, 11, 27)), (dt.date(2025, 12, 25), dt.date(2026, 11, 12))]
    sh = A.shift_shares(B, ordw)
    out["R0 ordinary-week fill by shift"] = {l: oh[l] * (1 - sh[l]) for l in B.locs}
    sh = A.shift_shares(B, [A.WINDOW])
    out["R1 last Christmas fill by shift"] = {l: oh[l] * (1 - sh[l]) for l in B.locs}
    models = {}
    for spec in ("noext", "noext_age", "noext_age_store", "sameday", "sameday_store", "golden", "ext_only", "ext_age",
                 "ext_dec"):
        models[spec] = C.fit(rows, spec, B.locs)
    for spec in ("noext", "noext_age", "noext_age_store", "sameday", "sameday_store"):
        out["R2 static " + spec] = C.Forward(B, models[spec]).static()
    for spec in ("noext_age", "sameday"):
        out["R2 cascade " + spec] = C.Forward(B, models[spec]).simulate(runs=8, seed=7)
    for spec in ("golden", "ext_only", "ext_age", "ext_dec"):
        out["R3 static " + spec] = C.Forward(B, models[spec]).static()
    for spec in ("ext_only", "ext_dec"):
        out["R3 cascade " + spec] = C.Forward(B, models[spec]).simulate(runs=8, seed=7)
    out.update(C.pool_routes(B, rows))
    return out, models


# ------------------------------------------------------------------------------------------------ the checks
def main(task, F):
    B, G, S, figs = F["B"], F["G"], F["S"], F["figs"]
    names = {l: B.loc_store[l] for l in B.locs}
    loc_of = {v: k for k, v in names.items()}
    CAMB, WERR = loc_of["Camberwell"], loc_of["Werribee"]
    target = B.target
    files = sorted(os.listdir(target))
    meta = json.load(open(os.path.join(task, "metadata.json"), encoding="utf-8"))

    # input gates
    check("G1 ten or more files", len(files) >= 10, "%d files" % len(files))
    check("G2 three or more formats", len({f.rsplit(".", 1)[1] for f in files}) >= 3,
          ", ".join(sorted({f.rsplit(".", 1)[1] for f in files})))
    check("G3 a file of 25,000 rows or more", len(B.offers) >= 25000, "%s %d rows" % (K.F_OFFERS, len(B.offers)))
    dist = meta["distractor_files"]
    check("G4 two distractors, named only in metadata.json", len(dist) >= 2 and all(d in files for d in dist)
          and not any("distract" in f.lower() for f in files), ", ".join(dist))
    check("G5 deliverables", all(os.path.exists(os.path.join(task, "golden", d)) for d in K.DELIVERABLES),
          ", ".join(K.DELIVERABLES))

    rows = G["rows"]
    R, models = routes(B, rows)

    # the three rules
    heads = open(os.path.join(target, K.F_OFFERS), encoding="utf-8").readline().strip().split(",")
    check("R1a the offer log ships at offer grain with each response", len(B.offers) > 100000 and "response" in heads
          and not any(w in ",".join(heads).lower() for w in ("extension", "abut", "forecast", "unworked")),
          "%d offers; columns %s" % (len(B.offers), ",".join(heads)))
    n_ext = sum(r["ext"] for r in rows)
    check("R1b the decisive term is built from the rosters, through the payroll key", n_ext > 1000 and
          not any("ext" in h for f in (K.F_REGISTER, K.F_KWS, K.F_LEGACY)
                  for h in open(os.path.join(target, f), encoding="utf-8").readline().lower().split(",")),
          "%d offers carry on from the casual's own shift" % n_ext)
    r0 = R["R0 ordinary-week register fill"]
    o0, m0 = C.rank(r0)
    check("R3a the default path files Werribee", o0[0] == WERR and m0 >= 1.15,
          "%s x%.2f; crew hours %.0f" % (names[o0[0]], m0, G["open"][o0[0]]))
    easy = [k for k in R if k.startswith(("R0", "R1", "R2", "P ", "open"))]
    worst = min(ratio(R[k], C.rank(R[k])[0][0], CAMB) for k in easy)
    check("R3a every easy route files another store, Camberwell 15 per cent behind", all(
        C.rank(R[k])[0][0] != CAMB for k in easy) and worst >= 1.15,
          "%d routes, narrowest leader over Camberwell x%.2f" % (len(easy), worst))
    dec = [k for k in R if k.startswith("R3")]
    dworst = min(C.rank(R[k])[1] for k in dec)
    check("R3b the offer model with only the trip-cost term added files Camberwell", all(
        C.rank(R[k])[0][0] == CAMB for k in dec) and dworst >= 1.2, "%d constructions, narrowest x%.2f" % (len(dec),
                                                                                                         dworst))
    gw, gm = G["order"][0], G["margin"]
    check("R3b the golden (cascade, 24 runs) files Camberwell", gw == CAMB and gm >= 1.2, "x%.2f over %s" % (
        gm, names[G["order"][1]]))
    wk = collections.defaultdict(lambda: collections.Counter())
    for ln in B.xmas:
        if ln["open"]:
            wk[ln["loc"]][ln["day"] // 7] += C.paid_min(ln["e"] - ln["s"])
    series = [tuple(sorted(v.items())) for v in wk.values()]
    lines = open(os.path.join(target, K.F_OFFERS), encoding="utf-8").read().splitlines()[1:]
    blocks = collections.Counter(tuple(x.split(",", 2)[2] for x in lines[i:i + 50]) for i in range(0, len(lines) - 50, 997))
    check("R3c no plant: no copied block, no memo, draft or stakeholder figure", len(set(series)) == len(series)
          and max(blocks.values()) == 1 and not any(re.search(r"memo|draft|note_to|request", f) for f in files),
          "%d open-hour series distinct; %d sampled offer blocks unique" % (len(series), len(blocks)))

    # Layer 1, each item mishandled: the call holds
    def static_call(rows2, B2=None):
        m = C.fit(rows2, "golden", B.locs)
        u = C.Forward(B2 or B, m).static()
        return C.rank(u)

    o, m = static_call(C.history(B, convert=False))
    check("R2.1 offer times read as store time: the call holds", o[0] == CAMB, "x%.2f" % m)
    cur = [dict(r, cl=B.people[r["pc"]]["home_location"] != r["loc"]) for r in rows]
    o, m = static_call(cur)
    check("R2.2 home store read current: the call holds", o[0] == CAMB, "x%.2f" % m)
    span = [dict(r, len_h=r["len_h"] + (0.5 if r["len_h"] > 5 else 0.0)) for r in rows]
    o, m = static_call(span)
    check("R2.3 span hours for paid: the call holds", o[0] == CAMB, "x%.2f" % m)
    B2 = copy.copy(B)
    B2.roster = [x for x in B.roster if x[1] >= C.dn(W.KWS_ROSTERS_FROM)]
    rows_k = C.history(B2)
    o, m = static_call(rows_k)
    check("R2.4 the legacy roster export left out: the call holds", o[0] == CAMB, "x%.2f" % m)
    B3 = copy.copy(B)
    B3.register = [r for r in B.register if r["outcome"] != "FILLED"]
    o, m = static_call(C.history(B3))
    check("R2.5 call-in fills left out of each casual's day: the call holds", o[0] == CAMB, "x%.2f" % m)
    B4 = copy.copy(B)
    C.identities(B4, follow_rehires=False)
    o, m = static_call(C.history(B4))
    check("R2.6 rehires' earlier payroll numbers not followed: the call holds", o[0] == CAMB, "x%.2f" % m)
    oh = C.open_hours(B)
    n18 = sum(C.paid_min(ln["e"] - ln["s"]) / 60.0 for ln in B.xmas if ln["open"] and ln["e"] == 1440)
    check("R2.7 shifts ending 24:00 only at Werribee; the crew hours do not touch them", n18 > 0 and all(
        ln["loc"] == WERR for ln in B.xmas if ln["e"] == 1440), "%.0f open hours end at midnight" % n18)
    rows_n = [dict(r, cl=C.home_loc_newest(B, r["pc"], C.D0 + dt.timedelta(days=r["day"])) != r["loc"]) for r in rows]
    o, m = static_call(rows_n)
    check("R2.8 re-engagements' cost centres read on the rows' own dates: the call holds", o[0] == CAMB, "x%.2f" % m)
    o, m = static_call(C.history(B, once=False))
    check("R2.9 a duplicate call-in's FILLED close kept in the casual's day: the call holds", o[0] == CAMB, "x%.2f" % m)

    # the ladder
    o1, m1 = C.rank(R["R1 last Christmas register fill"])
    check("L1 rung 0 Werribee, Camberwell fourth or lower", o0[0] == WERR and o0.index(CAMB) >= 3,
          "Camberwell %d" % (o0.index(CAMB) + 1))
    check("L2 rung 1 (last Christmas's fill) Werribee", o1[0] == WERR and m1 >= 1.15 and o1.index(CAMB) >= 2,
          "x%.2f, Camberwell %d" % (m1, o1.index(CAMB) + 1))
    r2 = [k for k in R if k.startswith("R2")]
    seconds = sum(1 for k in r2 if C.rank(R[k])[0].index(CAMB) == 1)
    check("L3 rung 2 (offer model, no trip cost) Werribee on every variant", all(C.rank(R[k])[0][0] == WERR for k in r2),
          "%d variants; Camberwell second on %d" % (len(r2), seconds))
    sig = np.sqrt(np.diag(models["golden"]["cov"]))
    mods = []
    j_ext = len(models["golden"]["b"]) - 2          # the golden's last two terms: the carry-on, a separate second shift
    for sgn in (-1, 1):
        mm = dict(models["golden"])
        b = mm["b"].copy()
        b[j_ext] += sgn * sig[j_ext]
        mm["b"] = b
        mods.append(C.rank(C.Forward(B, mm).static()))
    check("L4 a one-sigma move in the trip-cost term keeps Camberwell", all(o[0] == CAMB for o, _ in mods),
          "x%.2f and x%.2f" % (mods[0][1], mods[1][1]))
    Wd, truth, med, runs = world_truth(32)
    ot, mt = C.rank(truth)
    check("L5 the world's own Christmas (32 runs) puts Camberwell first", ot[0] == CAMB and mt >= 1.3,
          "x%.2f over %s" % (mt, names[ot[1]]))
    check("L6 the golden forecast and the world agree on the leader", ot[0] == G["order"][0])

    # semantic forks
    days = sorted({ln["day"] for ln in B.xmas})
    check("SF1 every Christmas line falls in 27 Nov to 24 Dec 2026, none on Christmas Day",
          C.D0 + dt.timedelta(days=days[0]) == W.PEAK26[0] and C.D0 + dt.timedelta(days=days[-1]) == W.PEAK26[1]
          and all(C.D0 + dt.timedelta(days=d) != dt.date(2026, 12, 25) for d in days))
    check("SF2 the crew hours are the store's open paid hours, not its unworked hours",
          figs["crew_hours"] == int(round(G["open"][CAMB] / 10.0)) * 10 and
          int(round(G["unworked"][CAMB] / 10.0)) * 10 != figs["crew_hours"],
          "%s against %.0f" % (figs["crew_hours"], G["unworked"][CAMB]))
    span_c = sum((ln["e"] - ln["s"]) / 60.0 for ln in B.xmas if ln["open"] and ln["loc"] == CAMB)
    check("SF3 span hours give another crew-hours bin, same store", int(round(span_c / 10.0)) * 10 != figs["crew_hours"],
          "span %.0f against paid %.1f" % (span_c, G["open"][CAMB]))
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(target, K.F_XMAS), read_only=True)
    bad = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=7, values_only=True):
            if row[0] is None:
                continue
            bad += (row[8] == "OPEN") != (row[6] is None)
    check("SF4 a line is OPEN exactly when no employee is on it", bad == 0)
    ru = G["order"][1]
    fw = C.Forward(B, G["model"])
    fw.open = [ln for ln in fw.open if ln["loc"] != ru]
    u5 = fw.simulate(runs=8, seed=11)
    u5 = {l: v for l, v in u5.items() if l != ru}
    check("SF5 the crew at the runner-up instead of nowhere: same call", C.rank(u5)[0][0] == CAMB, "x%.2f" % C.rank(u5)[1])
    regloc = {}
    for r in B.register:
        regloc[r["callin_id"]] = C.utc_to_local_min(r["shift_start_utc"])
    mism = 0
    for o_ in B.offers[::37]:
        d = dt.date.fromisoformat(o_["shift_date"])
        mism += abs(regloc[o_["callin_id"]] - (C.dn(d) * 1440 + C.hm(o_["shift_start"]))) > 0.01
    check("SF6 the offer log's store-time shifts agree with the register's UTC shifts", mism == 0)
    u7 = {l: v for l, v in G["unworked"].items() if names[l] != "Ballarat"}
    check("SF7 Ballarat left out: same call", C.rank(u7)[0][0] == CAMB)
    check("SF8 one Christmas roster version ships", sum(1 for f in files if "christmas_roster" in f) == 1)
    check("SF9 the median over the world's runs gives the same leader", C.rank(med)[0][0] == CAMB)
    feas = crew_feasible(B)
    check("SF10 25 crew under the EA limits can work every open shift at any store", all(feas.values()),
          ", ".join("%s %s" % (names[l], "yes" if v else "NO") for l, v in feas.items()))
    agree, seen = order_agrees(B)
    check("SF11 the offer log's order puts a store's own casuals before sister-store ones", agree == seen and seen > 500,
          "%d of %d postings" % (agree, seen))

    # closure axes not covered above
    check("C8 a trip-cost model fitted on ordinary weeks alone keeps Camberwell",
          C.rank(C.Forward(B, C.fit(rows, "golden", B.locs, mask=[not r["dec"] for r in rows])).static())[0][0] == CAMB)
    na = [dict(r, y=r["y"]) for r in rows if r["answer"] != "NO RESPONSE"]
    check("C19 NO RESPONSE left out of the fit: same call", C.rank(C.Forward(B, C.fit(na, "golden", B.locs)).static())[0][0]
          == CAMB)
    perm = list(range(len(rows)))
    np.random.default_rng(3).shuffle(perm)
    check("C14 row order does not matter", np.allclose(C.fit([rows[i] for i in perm], "golden", B.locs)["b"],
                                                       models["golden"]["b"], atol=1e-6))
    check("C22 the Christmas roster covers every store and every trading day", all(
        len({ln["day"] for ln in B.xmas if ln["loc"] == l}) == 28 for l in B.locs))
    C.index_availability(B)
    code = {c.idx: pc for c, pc in ((c, pc) for c in Wd.cas for pc in [None])}
    import records as RR
    RR.keys(Wd)
    mism = sum(1 for c in Wd.cas if c.active(W.AS_OF) for n in range(C.dn(W.PEAK26[0]), C.dn(W.PEAK26[1]) + 1)
               if [tuple(x) for x in C.windows_on(B, Wd.code[c.idx], C.D0 + dt.timedelta(days=n))]
               != [tuple(x) for x in W.windows(c, C.D0 + dt.timedelta(days=n))])
    check("C22b the profile in force at the extract gives every casual's windows on every Christmas date", mism == 0,
          "%d casual-days differ" % mism)
    ch = G["crew_hours"]
    edge = 5.0 - abs(ch - round(ch, -1))
    check("C10 the crew hours sit inside their nearest-10 bin", edge > 1e-9, "%.1f, %.1f from the bin edge" % (ch, edge))

    # the asks
    def figs_a(d):
        return [d[l][k] for l in B.locs for k in ("put_out", "covered", "uncovered")]

    def figs_b(d):
        return [d[l][k] for l in B.locs for k in ("imports", "exports")]

    ga, gb = figs_a(S["A_golden"]), figs_b(S["B_golden"])

    def off_a(name):
        return sum(abs(round(x, 1) - round(y, 1)) >= 0.05 for x, y in zip(figs_a(S[name]), ga))

    def off_b(name):
        return sum(abs(round(x, 1) - round(y, 1)) >= 0.05 for x, y in zip(figs_b(S[name]), gb))

    # the world's own last Christmas: each line put out once, covered when its last fill held, the filler's store
    tA, tB = world_asks(Wd, B)
    check("ASK the golden tables equal the world's own last Christmas", all(
        abs(tA[l][k] - S["A_golden"][l][k]) < 1e-6 for l in B.locs for k in ("put_out", "covered", "uncovered")) and all(
        abs(tB[l][k] - S["B_golden"][l][k]) < 1e-6 for l in B.locs for k in ("imports", "exports")))
    # the readings the module notes, the transmittal and the prompt file, each run as the alternative
    for name, label, need in (("A_S1_rows", "every call-in a shift, re-raises counted", 18),
                              ("A_S1b_noninstant", "every call-in a shift but those closed the minute raised", 18),
                              ("A_S3_nodrops", "a given-back shift's re-raise counted as another shift, its first "
                                               "fill kept (the round-2 judge's Fork S1)", 6),
                              ("A_S3c_keepfill", "a given-back shift counted once, its first fill kept as the cover",
                               4),
                              ("A_S4_utc", "shift dates read in UTC", 6),
                              ("A_S5_span", "span hours for paid", 18),
                              ("A_S6_noopn", "open lines left out of the hours put out", 16)):
        check("ASK A filed: %s misses" % label, off_a(name) >= need, "%d of 24 figures off" % off_a(name))
    check("ASK B filed: the home store read current misses", off_b("B_S1_current") >= 12,
          "%d of 16 figures off" % off_b("B_S1_current"))
    check("ASK B filed: a given-back shift credited to its first filler misses", off_b("B_S3c_keepfill") >= 1,
          "%d of 16 figures off" % off_b("B_S3c_keepfill"))
    check("ASK B filed: every FILLED call-in counted as written misses (the round-2 judge's Fork S2)",
          off_b("B_S3_rows") >= 6, "%d of 16 figures off" % off_b("B_S3_rows"))
    check("ASK B hazard S2_nodrops", off_b("B_S2_nodrops") >= 1, "%d of 16 figures off" % off_b("B_S2_nodrops"))
    # a give-back on a call-in raised for an absence (the round-2 judge's Sunshine 5.5 hours): the notes' give-back
    # rule covers every call-in, so looking for give-backs only on open-line fills misses
    gb_abs = 0
    slot_abs = collections.defaultdict(list)
    regw = [x for x in A._reg(B) if A.WINDOW[0] <= x["d"] <= A.WINDOW[1]]
    for y in regw:
        if y["reason"] == "ABS":
            slot_abs[(y["loc"], y["day"], y["s"], y["e"])].append(y)
    for x in regw:
        if x["reason"] == "ABS" and x["outcome"] == "FILLED":
            p_ = str(C.payroll_on(B, x["by"], x["d"]))
            gb_abs += any(y["absent"] == p_ and y["raised"] > x["closed"]
                          for y in slot_abs[(x["loc"], x["day"], x["s"], x["e"])])
    check("ASK A filed: give-backs looked for only on open-line fills miss (fills of absence call-ins given back too)",
          off_a("A_S3d_opngb") >= 1 and gb_abs >= 1, "%d of 24 figures off; %d absence call-ins given back" % (
              off_a("A_S3d_opngb"), gb_abs))
    # the silent devices, each found in the data once the filing is read
    check("ASK A device A1: the stores' identical open lines merged with their re-raises misses", off_a("A_S2_keys") >= 18,
          "%d of 24 figures off" % off_a("A_S2_keys"))
    check("ASK A device A3: an absence's second report and the import's second raise each counted a shift misses",
          off_a("A_S3b_raise") >= 12, "%d of 24 figures off" % off_a("A_S3b_raise"))
    check("ASK A device A4: the 07:00 chain, the duplicate reports merged as filed, counts the import's second raise "
          "of an open line as a shift", off_a("A_chain") >= 18, "%d of 24 figures off" % off_a("A_chain"))
    cov_off = sum(abs(round(S["A_chain"][l]["covered"], 1) - round(S["A_golden"][l]["covered"], 1)) >= 0.05
                  for l in B.locs)
    check("ASK A device A4: the chain counts a twice-raised line's FILLED close as more cover", cov_off >= 6,
          "%d of 8 covered figures off" % cov_off)
    # the module offers a casual who gave a shift back nothing else that day, so voiding by the day's absence and
    # by the shift's own re-raise agree
    check("ASK the day's absence voids the same fills as the shift's re-raise", off_a("A_S7_dayvoid") == 0 and
          off_b("B_S7_dayvoid") == 0)
    # Ask B's own device: the new-year store moves payroll keyed as re-engagements, the new cost centre dated from
    # the keying, read from the number KWS holds
    check("ASK B device B2: re-engagements read on the cost-centre rows' own dates miss (golden shifts)",
          off_b("B_S8_newest") >= 9, "%d of 16 figures off" % off_b("B_S8_newest"))
    check("ASK B device: the shift-identity fold misses Table 2 with it", off_b("B_crack") >= 9,
          "%d of 16 off (%d with the golden's home read)" % (off_b("B_crack"), off_b("B_crack_payroll")))
    check("ASK B device: the 07:00 chain misses Table 2", off_b("B_chain") >= 9, "%d of 16 off (%d with the golden's "
          "home read)" % (off_b("B_chain"), off_b("B_chain_payroll")))
    # the cracker sheet, measured on the cut pack (generator/cracker.py)
    import cracker as CR
    sheet, blind = CR.measure(task)
    check("ASK each cracker path keeps few enough figures for a pair at or under 40", all(
        r["pair"] <= 40.0 for r in sheet.values() if not r["record"]), "; ".join(
        "%s %d of 40, pair %.1f" % (k, r["kept"], r["pair"]) for k, r in sheet.items()))
    check("ASK the round-2 judge's own route, run for the record, also sits at or under 40", all(
        r["pair"] <= 40.0 for r in sheet.values() if r["record"]))
    check("ASK the blind responses recover no more than a quarter of either table", max(
        blind["rate_A"], blind["nil_A"], blind["const_A"]) <= 6 and max(blind["nil_B"], blind["const_B"]) <= 4,
          str(blind))
    exact = all(abs(x * 2 - round(x * 2)) < 1e-6 for x in ga + gb)
    check("ASK every figure is exact at one decimal (paid hours come in half hours)", exact)
    pub = collections.Counter()
    rep = collections.Counter()
    for x in A._reg(B):
        if A.WINDOW[0] <= x["d"] <= A.WINDOW[1]:
            if A._is_repost(x):
                rep[x["loc"]] += 1
    check("ASK A1 reposts sit in the window at every store with call-ins", all(rep[l] > 0 for l in B.locs),
          ", ".join("%s %d" % (names[l], rep[l]) for l in B.locs))
    vac = collections.Counter()
    for loc, day, s_, e_, pn, st in B.roster:
        d = C.D0 + dt.timedelta(days=day)
        if st == "VACANT" and A.WINDOW[0] <= d <= A.WINDOW[1]:
            vac[loc] += 1
    orig = collections.Counter(x["loc"] for x in A._reg(B) if A.WINDOW[0] <= x["d"] <= A.WINDOW[1]
                               and x["reason"] == "OPN" and x["r"]["raised_at"] == A.PUBLISHED_25)
    check("ASK A4 the 2025 import raised some open lines twice: postings at publication cover every vacant line and "
          "exceed them at six or more stores", all(orig[l] >= vac[l] for l in B.locs) and
          sum(orig[l] > vac[l] for l in B.locs) >= 6, ", ".join("%s %d of %d" % (names[l], orig[l], vac[l])
                                                                for l in B.locs))
    movers = [pc for pc in B.people if B.people[pc]["home_location"] != C.home_loc(B, pc, dt.date(2025, 12, 10))]
    check("ASK B1 store moves touch every store", len({B.people[pc]["home_location"] for pc in movers} |
                                                     {C.home_loc(B, pc, dt.date(2025, 12, 10)) for pc in movers})
          == len(B.locs), "%d casuals moved" % len(movers))

    # separation: the asks' devices never move the call
    check("SEP1 the golden reads no register outcome but fills and vacated lines",
          C.rank(C.Forward(B, C.fit(C.history(B3), "golden", B.locs)).static())[0][0] == CAMB)
    check("SEP2 the asks read neither the Christmas roster nor the offer log",
          "xmas" not in open(os.path.join(os.path.dirname(__file__), "asks.py")).read().replace("XMAS", "")
          and "B.offers" not in open(os.path.join(os.path.dirname(__file__), "asks.py")).read())

    # realism
    text = ""
    for sec in (DOC.STANDARD, DOC.NOTES, DOC.EA):
        for t, ps in sec:
            text += t + " " + " ".join(ps)
    text += open(os.path.join(target, K.F_TRANSMIT), encoding="utf-8").read()
    check("RZ1 no em or en dash in any shipped text", "\u2014" not in text and "\u2013" not in text)
    card = json.load(open(os.path.join(os.path.dirname(task), ".claude", "skills", "fingerprint", "cards",
                                       "task155.json"), encoding="utf-8"))
    people = set(card["world"]["people"])
    named = set(re.findall(r"(Beth Hill|Meredith Stevens|Kristina Patel|Belinda Friedman|[A-Z][a-z]+ [A-Z][a-z]+, "
                           r"(?:Workforce|General|Payroll))", text))
    check("RZ2 every person named is a drawn persona", all(n.split(",")[0] in people for n in named), ", ".join(sorted(named)))
    dates = [dt.date(int(y), int(m), int(d)) for d, m, y in re.findall(r"\b(\d{1,2})/(\d{1,2})/(20\d\d)\b", text)]
    check("RZ3 no document dated after the as-of", all(d <= W.AS_OF for d in dates), "latest %s" % max(dates))
    listed = set(re.findall(r"^(\S+\.(?:csv|xlsx|json|pdf|txt))\s", open(os.path.join(target, K.F_TRANSMIT),
                                                                         encoding="utf-8").read(), re.M))
    check("RZ7 the transmittal lists every shipped file but itself", listed == set(files) - {K.F_TRANSMIT},
          "%d listed" % len(listed))
    st = {os.path.getmtime(os.path.join(target, f)) for f in files}
    check("RZ4 every file carries the extract's stamp", len(st) == 1)
    tot = [round(G["open"][l], 1) for l in B.locs]
    check("RZ5 no round headline total", all(abs(x / 50.0 - round(x / 50.0)) > 1e-9 for x in tot), str(tot))
    wcount = collections.Counter((o_["location_id"], o_["shift_date"][:7]) for o_ in B.offers)
    check("RZ6 no two stores share a month's offer count", len(set(wcount.values())) > 0.9 * len(wcount))
    # behavioural cells a reviewer would cross-tab: none empty of take-up. The second-shift cell is held well under the
    # day-off rate on purpose (DESIGN_NOTE, Tried and rejected: a higher rate lends the pooled same-day model the
    # carry-on effect)
    cell = collections.defaultdict(lambda: [0, 0])
    for r_ in rows:
        k_ = "ext" if r_["ext"] else ("split" if r_["split"] else "off")
        cell[k_][0] += 1
        cell[k_][1] += r_["y"]
    rate = {k_: v[1] / v[0] for k_, v in cell.items()}
    check("RZ10 a second, separate shift on a rostered day is taken, well under the day-off rate", cell["split"][1] >= 30
          and rate["split"] < rate["off"] / 5, "%.2f%% of %d (%d taken), a day off %.1f%%" % (
              100 * rate["split"], cell["split"][0], cell["split"][1], 100 * rate["off"]))
    near = sum(r_["near"] for r_ in rows)
    check("RZ11 no offer leaves a casual under an hour to get from or to a shift at another store (the notes' rule)",
          near == 0, "%d of %d offers" % (near, len(rows)))
    # the brief's statements about the two stores, held against the record (house fix H6)
    C.index_availability(B)
    fwp = C.Forward(B, None)
    fwp.reset()
    east = {loc_of["Box Hill"], loc_of["Ringwood"]}
    sizes, members = [], set()
    for ln in fwp.open:
        if ln["loc"] == CAMB and ln["code"] == "C09":
            el = [pc for pc in fwp.order[CAMB] if fwp.eligible(pc, ln)]
            sizes.append(len(el))
            members.update(el)
    sister = [pc for pc in members if fwp.home[pc] in east]
    young = [pc for pc in sister if 18 <= fwp.age[pc] < 25]
    check("RZ8 plenty of casuals are free for Camberwell's 9-to-3 uplift, most of them Box Hill's and Ringwood's "
          "18-to-24s", sizes and statistics.median(sizes) >= 20 and len(sister) > len(members) / 2 and
          len(young) > len(sister) / 2, "median %.0f a shift; %d casuals, %d from Box Hill or Ringwood, %d of them "
          "18 to 24" % (statistics.median(sizes), len(members), len(sister), len(young)))
    wcc = [r["cost_centre"] for r in B.stores if r["kws_location"] == WERR][0]
    feb = {int(r["employee_no"]) for r in B.payroll if r["engagement_start"].startswith("2026-02")
           and r["cost_centre"] == wcc and not r["previous_employee_no"]}
    lines = [ln for ln in B.xmas if ln["loc"] == WERR and ln["pn"] in feb]
    on = sum(ln["code"] in ("W08", "A14") for ln in lines)
    check("RZ9 Werribee's February intake is rostered mostly on the 8-to-2 and 2-to-6 shifts", on > len(lines) / 2,
          "%d of %d lines" % (on, len(lines)))

    # distractors
    used = open(os.path.join(os.path.dirname(__file__), "calc.py")).read() + open(
        os.path.join(os.path.dirname(__file__), "asks.py")).read()
    check("D1 the golden and the asks read neither distractor", "F_PLAN" not in used and "F_CAL" not in used)
    import openpyxl as _ox
    wsp = _ox.load_workbook(os.path.join(target, K.F_PLAN), read_only=True).worksheets[0]
    pnotes = [row[-1] for row in wsp.iter_rows(min_row=6, values_only=True) if row[0]]
    staffing = re.compile(r"casual|staff|roster|call-in|crew|shift|team|hire|labour|labor", re.I)
    check("D2 the trading plan's notes are trading notes, several stores', none on staffing or call-ins",
          sum(1 for n in pnotes if n) >= 5 and not any(staffing.search(n or "") for n in pnotes),
          "; ".join(n for n in pnotes if n))
    xcore = __import__("zipfile").ZipFile(os.path.join(target, K.F_XMAS)).read("docProps/core.xml").decode()
    stamps = re.findall(r"<dcterms:(created|modified)[^>]*>([^<]+)<", xcore)
    check("RZ12 the published roster workbook was saved by its 10:00 publication", all(
        v <= "2026-11-11T23:00:00Z" for _, v in stamps) and len(stamps) == 2, ", ".join("%s %s" % kv for kv in stamps))

    fails = [r for r in RESULTS if not r[1]]
    print("checks: %d, failed %d" % (len(RESULTS), len(fails)))
    _write_routes(task, R, names, truth, G)
    return 0 if not fails else 1


def crew_feasible(B):
    """a greedy roster of 25 crew: each day's open shifts by start, each to a crew member free then, under 11 paid
    hours a day, 38 a week and 10 hours between days"""
    out = {}
    for loc in B.locs:
        shifts = sorted([ln for ln in B.xmas if ln["open"] and ln["loc"] == loc], key=lambda l: (l["day"], l["s"], l["e"]))
        book = collections.defaultdict(list)
        wk = collections.Counter()
        ok = True
        for ln in shifts:
            placed = False
            for m in sorted(range(25), key=lambda m: (wk[(m, ln["day"] // 7)], m)):
                today = book[(m, ln["day"])]
                if any(s < ln["e"] and ln["s"] < e for s, e in today):
                    continue
                if sum(C.paid_min(e - s) for s, e in today) + C.paid_min(ln["e"] - ln["s"]) > 660:
                    continue
                if wk[(m, ln["day"] // 7)] + C.paid_min(ln["e"] - ln["s"]) / 60.0 > 38:
                    continue
                if any((ln["day"] * 1440 + ln["s"]) - ((ln["day"] - 1) * 1440 + e) < 600 for s, e in book[(m, ln["day"] - 1)]):
                    continue
                book[(m, ln["day"])].append((ln["s"], ln["e"]))
                wk[(m, ln["day"] // 7)] += C.paid_min(ln["e"] - ln["s"]) / 60.0
                placed = True
                break
            if not placed:
                ok = False
                break
        out[loc] = ok
    return out


def order_agrees(B):
    """for postings offered to both a store's own casual and a sister-store casual, every own offer came first"""
    first = {}
    seen = agree = 0
    by = collections.defaultdict(list)
    for o in B.offers:
        by[o["callin_id"]].append(o)
    for cid, lst in by.items():
        loc = lst[0]["location_id"]
        d = dt.date.fromisoformat(lst[0]["shift_date"])
        kinds = [C.home_loc(B, o["person_code"], d) == loc for o in lst]
        if all(kinds) or not any(kinds):
            continue
        seen += 1
        last_own = max(i for i, k in enumerate(kinds) if k)
        first_oth = min(i for i, k in enumerate(kinds) if not k)
        agree += last_own < first_oth
    return agree, seen


def _write_routes(task, R, names, truth, G):
    """the route table for the design note (generator/routes.txt, evidence, not shipped)"""
    lines = []
    for k, u in list(R.items()) + [("world truth (32 runs)", truth), ("golden cascade (24 runs)", G["unworked"])]:
        o, m = C.rank(u)
        if k.startswith("P "):
            # a pool route: the thinnest store leads; the pool measure itself is printed, casuals per shift or per hour
            lines.append("%-40s %-11s x%.2f | %s" % (k, names[o[0]], m, " ".join("%s %.1f" % (names[l], 1.0 / u[l])
                                                                                   for l in o)))
            continue
        lines.append("%-40s %-11s x%.2f | %s" % (k, names[o[0]], m, " ".join("%s %.0f" % (names[l], u[l]) for l in o)))
    open(os.path.join(task, "generator", "routes.txt"), "w").write("\n".join(lines) + "\n")
