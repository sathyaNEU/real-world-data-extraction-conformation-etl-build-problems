"""The bundle's workbooks and documents: Merriden's delisting workbook, data dictionary and corporate actions feed
specification; Ledgerpoint's user guide; Northcote's methodology summary and reconstitution workbook; Kelloway's
commission schedule; and the research data team's README for the extracts."""
import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import letter                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import inch                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import universe as U  # noqa: E402
import world as W     # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("LiberationSans", "Regular"), ("LiberationSans-Bold", "Bold"), ("LiberationSerif", "Regular"),
              ("LiberationSerif-Bold", "Bold")):
    try:
        pdfmetrics.getFont(nm)
    except KeyError:
        pdfmetrics.registerFont(TTFont(nm, FONT_DIR + nm.split("-")[0] + "-" + f + ".ttf"))
INK = colors.HexColor("#1b1d20")
GREY = colors.HexColor("#5b6068")


def styles(font, accent, size=9.4):
    bold = font + "-Bold"
    return dict(
        BODY=ParagraphStyle("b", fontName=font, fontSize=size, leading=size * 1.36, textColor=INK, spaceAfter=5),
        CELL=ParagraphStyle("c", fontName=font, fontSize=7.9, leading=9.8, textColor=INK),
        CELLB=ParagraphStyle("cb", fontName=bold, fontSize=7.9, leading=9.8, textColor=INK),
        SMALL=ParagraphStyle("s", fontName=font, fontSize=7.6, leading=9.8, textColor=GREY, spaceAfter=4),
        H1=ParagraphStyle("h1", fontName=bold, fontSize=14, leading=18, textColor=accent, spaceAfter=2),
        H1S=ParagraphStyle("h1s", fontName=font, fontSize=8.8, leading=12, textColor=GREY, spaceAfter=10),
        H2=ParagraphStyle("h2", fontName=bold, fontSize=10.2, leading=13.5, textColor=INK, spaceBefore=8,
                          spaceAfter=3))


def _doc(path, title, footer, author, st_font="LiberationSans"):
    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(GREY)
        canvas.setFont(st_font, 7.2)
        canvas.drawString(0.85 * inch, 0.55 * inch, footer)
        canvas.drawRightString(letter[0] - 0.85 * inch, 0.55 * inch, "%d" % doc.page)
        canvas.restoreState()
    d = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.85 * inch, rightMargin=0.85 * inch,
                          topMargin=0.75 * inch, bottomMargin=0.85 * inch, title=title, author=author, subject="",
                          creator="", producer="")
    return d, frame


def _table(rows, widths, st, head):
    data = [[Paragraph(str(v), st["CELLB"] if i == 0 else st["CELL"]) for v in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#a3a8ae")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(head)),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 3.5),
                           ("RIGHTPADDING", (0, 0), (-1, -1), 3.5)]))
    return t


def _story(title, sub, sections, st):
    out = [Paragraph(title, st["H1"]), Paragraph(sub, st["H1S"])]
    for head, paras in sections:
        if head:
            out.append(Paragraph(head, st["H2"]))
        for p in paras:
            out.append(Paragraph(p, st["BODY"]) if isinstance(p, str) else p)
    return out


# ------------------------------------------------------------------------------------ Merriden's data dictionary
def write_dictionary(path):
    st = styles("LiberationSans", colors.HexColor("#123c69"))
    d, frame = _doc(path, "Merriden End-of-Day US Equity: Data Dictionary 4.0",
                    "Merriden Market Data  |  End-of-Day US Equity  |  Data Dictionary 4.0  |  January 2024", W.VENDOR_PX)
    old = [["Field", "Type", "Description"],
           ["DATE", "integer", "Trading date, yyyymmdd."],
           ["SID", "integer", "Merriden security identifier. Permanent; never reassigned."],
           ["TICKER", "text", "Exchange ticker in force on DATE. Tickers are reassigned by the exchanges and can "
                              "refer to different securities on different dates."],
           ["EXCH", "integer", "Listing exchange on DATE: 1 NYSE, 2 NYSE American, 3 Nasdaq."],
           ["OPEN, HIGH, LOW", "decimal", "Consolidated opening, high and low trade prices. Blank when the security did "
                                         "not trade."],
           ["CLOSE", "decimal", "Consolidated closing price. When the security did not trade, the average of the "
                               "closing bid and ask (CLOSE_SRC M)."],
           ["CLOSE_SRC", "text", "T closing trade; M bid/ask average."],
           ["VOLUME", "integer", "Consolidated volume, shares, as first published."],
           ["BID, ASK", "decimal", "Closing inside quote."],
           ["SHROUT", "integer", "Shares outstanding, thousands."],
           ["ADJ_FACTOR", "decimal", "Multiply a price on DATE by ADJ_FACTOR to express it in shares as of the last "
                                     "date in the file (29 December 2023), or the security's last row if earlier. "
                                     "Divide a share quantity by it."],
           ["RET", "decimal", "Total return from the previous row's CLOSE: cash distributions with an ex-date on "
                              "DATE and the value of securities distributed (at their first regular-way close) are "
                              "added; splits are applied. Blank on a security's first row."],
           ["RETX", "decimal", "As RET with no distributions added."]]
    new = [["Column", "Type", "Description"],
           ["trade_date", "date", "Trading date."],
           ["security_id", "int64", "As SID."],
           ["symbol", "string", "Exchange ticker in force on trade_date."],
           ["mic", "string", "Listing exchange MIC on trade_date: XNYS, XASE, XNAS."],
           ["open_px, high_px, low_px", "float64", "Consolidated opening, high and low trade prices; null when the "
                                                   "security did not trade."],
           ["close_px", "float64", "Official consolidated close; bid/ask average when there was no trade."],
           ["close_type", "string", "TRADE or MID."],
           ["volume", "int64", "Consolidated volume, shares, including trades reported late."],
           ["closing_bid, closing_ask", "float64", "Closing inside quote."],
           ["shares_outstanding", "int64", "Shares outstanding, units."],
           ["px_adj_factor", "float64", "Multiply a price on trade_date by px_adj_factor to express it in shares as of "
                                        "the delivery's last trading date, or the security's last row if earlier."],
           ["total_return", "float64", "As RET."],
           ["price_return", "float64", "As RETX."]]
    secs = [
        ("1. Deliveries", [
            "The legacy platform's history file covers 2 January 2018 to 29 December 2023 and is final. The current "
            "platform delivers Parquet from its first production date, 2 January 2024, with a backfill from 2 October "
            "2023 so that the two deliveries overlap for one quarter. In the overlap the current platform carries the "
            "official closes and volumes after the consolidated tape's late corrections; the legacy file carries them "
            "as first published. Values in the overlap are not reconciled between the two deliveries."]),
        ("2. Legacy layout (CSV)", [_table(old, [1.15 * inch, 0.65 * inch, 4.9 * inch], st, "#e2e8f0")]),
        ("3. Current layout (Parquet)", [_table(new, [1.55 * inch, 0.65 * inch, 4.5 * inch], st, "#e2e8f0")]),
        ("4. Conventions", [
            "Rows. A security has a row for every trading day from its first regular-way trade to its last exchange "
            "trade, including days it did not trade. Rows end at delisting; values after delisting, including the "
            "first over-the-counter price and the delisting return, are in the delisting workbook.",
            "Adjustment. Prices are as traded. A split or reverse split changes the price on its ex-date; the price "
            "factor and the returns carry it. The two deliveries anchor their factors to different dates, so factors "
            "from the two files are not interchangeable.",
            "Shares outstanding are updated from each filing on the filing date and adjusted on a split's ex-date.",
            "Identifiers. SID and security_id are the same number for the same security. A ticker or a CUSIP can "
            "change during a security's life, and a ticker can later be issued to another company; join on the "
            "identifier and the date, using the security master."]),
    ]
    story = _story("End-of-Day US Equity: Data Dictionary", "Version 4.0, January 2024. Merriden Market Data, "
                   "Reference Data Products.", secs, st)
    d.build(story, onFirstPage=frame, onLaterPages=frame)


# ------------------------------------------------------------------------------------ the CA feed specification
def write_ca_spec(path):
    st = styles("LiberationSans", colors.HexColor("#123c69"))
    d, frame = _doc(path, "Merriden Corporate Actions Feed 3.2: Specification",
                    "Merriden Market Data  |  Corporate Actions Feed 3.2  |  June 2023", W.VENDOR_PX)
    types = [["event_type", "terms"],
             ["DIV_CASH, DIV_SPECIAL", "gross_amount (per share, before withholding), currency, frequency"],
             ["SPLIT, REVERSE_SPLIT", "new_shares, old_shares (new_shares for each old_shares held)"],
             ["SPINOFF", "distributed_security_id, distributed_symbol, distributed_cusip, ratio (distributed shares "
                         "per share held)"],
             ["MERGER_CASH, GO_PRIVATE", "acquirer_name, cash_per_share"],
             ["MERGER_STOCK", "acquirer_name, acquirer_security_id, acquirer_symbol, shares_per_share, "
                              "cash_per_share"],
             ["LISTING", "listing_type (IPO, SPIN-OFF), offer_price for an IPO"],
             ["DELISTING", "reason, successor_venue"],
             ["SYMBOL_CHANGE, NAME_CHANGE, EXCHANGE_CHANGE", "old_value, new_value"]]
    fields = [["Field", "Description"],
              ["ca_id", "Event identifier. All revisions of an event share it."],
              ["revision", "1 for the first record of an event, incremented with each later record."],
              ["status", "NEW; AMEND (replaces every earlier revision of the event); CANCEL (the event will not take "
                         "place); COMPLETE (a merger or acquisition has closed: effective_date is the first day the "
                         "target no longer trades)."],
              ["published_utc", "When the revision was published, UTC."],
              ["security_id, symbol, cusip", "The security, with the symbol and CUSIP in force at announcement."],
              ["announce_date", "The issuer's announcement date."],
              ["ex_date, record_date, pay_date", "As announced. The record date follows the settlement cycle in force "
                                                 "on the ex-date."],
              ["effective_date", "For listings, delistings, completed mergers and changes of symbol, name or "
                                 "exchange."],
              ["terms", "Event terms, by event type (section 3)."],
              ["note", "Free text where the issuer's notice gives a reason."]]
    secs = [
        ("1. Delivery", ["One JSON document per extract with a header and a records array. Every revision of every "
                         "event published up to the extract time is included; a record is never edited after "
                         "publication. Events announced before the extract can have dates after it."]),
        ("2. Record fields", [_table(fields, [1.6 * inch, 5.1 * inch], st, "#e2e8f0")]),
        ("3. Event types", [_table(types, [2.3 * inch, 4.4 * inch], st, "#e2e8f0")]),
    ]
    story = _story("Corporate Actions Feed 3.2", "Specification, June 2023. Merriden Market Data.", secs, st)
    d.build(story, onFirstPage=frame, onLaterPages=frame)


# ------------------------------------------------------------------------------------ Ledgerpoint's guide
def write_fund_guide(path):
    st = styles("LiberationSerif", colors.HexColor("#3b2f63"), 9.6)
    d, frame = _doc(path, "Ledgerpoint Point-in-Time Fundamentals: User Guide",
                    "Ledgerpoint Financial Data  |  Point-in-Time Quarterly Fundamentals  |  User Guide, March 2025",
                    W.VENDOR_FUND, "LiberationSerif")
    srcs = [["source", "What the row holds", "knowledge_date"],
            ["PRESS_RELEASE", "The results announcement: income statement and per-share figures for the quarter "
                              "(a fourth quarter's own figures included); balance sheet items where the release has "
                              "them; no cash flow items.", "Date of the release."],
            ["10-Q", "The quarterly report: income statement for the quarter, balance sheet at period end, cash flow "
                     "items for the fiscal year to date.", "Filing date."],
            ["10-K", "The annual report: income statement for the fiscal year (period_type FY, fiscal_quarter "
                     "blank), balance sheet at year end, cash flow items for the year.", "Filing date."],
            ["10-Q/A, 10-K/A", "An amended report restating the period. It replaces the earlier rows from its "
                               "knowledge_date.", "Filing date of the amendment."]]
    cal = [["Field", "Description"],
           ["lp_id", "Ledgerpoint entity identifier."],
           ["ticker, company", "As of the event."],
           ["fiscal_period", "Fiscal year and quarter, company convention."],
           ["event_utc", "Scheduled time of the release, UTC. PROJECTED rows carry a date with no time (midnight "
                         "New York time)."],
           ["time_code", "BMO before the market opens, AMC after it closes, DMT during market hours, UNK unknown."],
           ["status", "PROJECTED (Ledgerpoint's estimate) or CONFIRMED (from the company)."],
           ["updated_utc", "When the row was written. For a fiscal period, the row with the latest updated_utc "
                           "supersedes the others."]]
    secs = [
        ("1. Point in time", [
            "Every row is stamped with the date its figures became available (knowledge_date). A period appears "
            "once for each source that reported it. Using only rows with a knowledge_date on or before a date "
            "reproduces what was known that day. Figures are never revised in place."]),
        ("2. Sources", [_table(srcs, [1.1 * inch, 4.0 * inch, 1.6 * inch], st, "#ebe7f3")]),
        ("3. Units and per-share figures", [
            "Monetary amounts are in the unit the company reports in, given in units: thousands for most companies, "
            "millions for others. shares_diluted is in thousands for every company. eps_diluted is in dollars as "
            "reported at the time; per-share figures are not restated for later splits. Cash flow items in 10-Q "
            "rows are for the fiscal year to date."]),
        ("4. Fiscal calendars", [
            "fiscal_year and fiscal_quarter follow each company's own convention. Most companies name a fiscal year "
            "after the calendar year in which it ends; companies on a 52-53 week year ending near the end of January "
            "name it after the calendar year in which it begins, and their periods end on a Saturday."]),
        ("5. Identifiers", [
            "lp_id is permanent. ticker and cusip are those in force on the knowledge_date; both can change."]),
        ("6. Earnings calendar", [_table(cal, [1.2 * inch, 5.5 * inch], st, "#ebe7f3")]),
    ]
    story = _story("Point-in-Time Quarterly Fundamentals", "User Guide, March 2025. Ledgerpoint Financial Data.",
                   secs, st)
    d.build(story, onFirstPage=frame, onLaterPages=frame)


# ------------------------------------------------------------------------------------ Northcote's methodology
def write_methodology(path):
    st = styles("LiberationSans", colors.HexColor("#0f5257"))
    d, frame = _doc(path, "Northcote US Small Cap Select Index: Methodology Summary",
                    "Northcote Indices  |  US Small Cap Select (NUSSC)  |  Methodology Summary, February 2025",
                    W.INDEX_CO)
    secs = [
        ("1. Eligibility", [
            "US-listed common stock with at least 60 trading days of listing on the rank day, a closing price of at "
            "least $1.00 and a total market capitalisation (all shares outstanding times the close) between $300 "
            "million and $7.5 billion on the rank day."]),
        ("2. Reconstitution", [
            "Rank day: the last trading day of May. The 80 largest eligible companies by total market capitalisation "
            "form the index from the close of the fourth Friday of June (the preceding trading day if that Friday is "
            "a holiday). The reconstitution workbook lists each June's constituents with their index shares and "
            "weights at that close."]),
        ("3. Weighting", [
            "Index shares are shares outstanding times a float factor, scaled at each reconstitution by a capping "
            "factor so that no constituent's weight exceeds 4 per cent at the reconstitution close. Index shares "
            "stay fixed between reconstitutions except that a split's ratio is applied on its ex-date."]),
        ("4. Changes between reconstitutions", [
            "A constituent leaves the index after the close of its last trading day. It is not replaced until the "
            "next reconstitution. Initial public offerings that meet the size limits on the quarterly addition date "
            "(the third Friday of March, September and December) join after that day's close. A security spun off "
            "by a constituent joins on its ex-date."]),
        ("5. Calculation", [
            "The price return index moves each day by the change in the index shares' value at the previous close, "
            "a constituent's own move including the value of anything it distributed other than cash. The total "
            "return index also includes cash dividends on their ex-dates. Levels are published to two decimals."]),
    ]
    story = _story("US Small Cap Select Index", "Methodology Summary, February 2025. Northcote Indices.", secs, st)
    d.build(story, onFirstPage=frame, onLaterPages=frame)


# ------------------------------------------------------------------------------------ Kelloway's schedule
def write_fees(path):
    st = styles("LiberationSans", colors.HexColor("#5c3d11"))
    d, frame = _doc(path, "Kelloway Securities: Commission and Fee Schedule",
                    "Kelloway Securities  |  Institutional Equities  |  Schedule KS-PBK-2025-1", W.BROKER)
    t = [["Item", "Rate", "Basis"],
         ["Commission, algorithmic execution (trades from 3 March 2025)", "$0.0100", "per share executed"],
         ["Commission, algorithmic execution (superseded schedule of 1 August 2022)", "$0.0150", "per share executed"],
         ["Clearing", "$0.00025", "per share executed, buys and sells"]]
    secs = [
        ("Client", ["Pellbrook Capital LLC, account PBK-SMID-01. Effective for trades from 3 March 2025; replaces "
                    "the schedule of 1 August 2022."]),
        ("Rates", [_table(t, [3.9 * inch, 0.9 * inch, 1.9 * inch], st, "#f1e9dc")]),
        ("Execution reports", [
            "Executions are reported per fill with the client order identifier. Times are New York local time. "
            "exec_type F is a fill; X cancels the fill named in ref_exec_id; C replaces the fill named in "
            "ref_exec_id with corrected terms. Venue is the market identifier code of the execution venue; KLWX is "
            "Kelloway's own alternative trading system. Settlement is on the standard cycle for the trade date."]),
    ]
    story = _story("Commission and Fee Schedule", "Kelloway Securities, Institutional Equities. Issued 10 February "
                   "2025.", secs, st)
    d.build(story, onFirstPage=frame, onLaterPages=frame)


# ------------------------------------------------------------------------------------ workbooks
HEAD = PatternFill("solid", fgColor="DCE3EA")


def _sheet(ws, header, rows, widths, fmts=None, freeze="A2"):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD
        c.alignment = Alignment(vertical="top", wrap_text=True)
    for r in rows:
        ws.append(r)
    for n, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(n)].width = w
    if fmts:
        for col, fmt in fmts.items():
            for cell in ws[col][1:]:
                cell.number_format = fmt
    ws.freeze_panes = freeze


def write_delistings(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Delistings"
    head = ["security_id", "last_symbol", "issuer_name", "cusip", "last_trading_date", "final_exchange_close", "code",
            "reason", "cash_per_share", "shares_per_share", "acquirer_security_id", "otc_first_date",
            "otc_first_price", "delisting_return", "note"]
    data = [[r["security_id"], r["last_symbol"], r["issuer_name"], r["cusip"], r["last_trading_date"],
             r["final_exchange_close"], r["code"], r["reason"],
             float(r["cash_per_share"]) if r["cash_per_share"] is not None else None,
             float(r["shares_per_share"]) if r["shares_per_share"] is not None else None, r["acquirer_security_id"],
             r["otc_first_date"], float(r["otc_first_price"]) if r["otc_first_price"] is not None else None,
             float(r["delisting_return"]) if r["delisting_return"] is not None else None, r["note"]] for r in rows]
    _sheet(ws, head, data, [11, 10, 38, 11, 13, 12, 6, 26, 12, 12, 12, 12, 12, 13, 40],
           {"E": "yyyy-mm-dd", "L": "yyyy-mm-dd", "F": "0.0000", "I": "0.00", "J": "0.0000", "M": "0.0000",
            "N": "0.000000"})
    ws2 = wb.create_sheet("Codes")
    _sheet(ws2, ["code", "meaning", "delisting_return"],
           [["M1", "Merger or acquisition, cash consideration", "cash_per_share over final_exchange_close, less one"],
            ["M2", "Merger or acquisition, stock consideration",
             "shares_per_share times the acquirer's close on the effective date, over final_exchange_close, less one"],
            ["P1", "Going-private transaction, cash consideration", "as M1"],
            ["X1", "Delisted by the exchange: continued listing standards", "otc_first_price over "
                                                                             "final_exchange_close, less one"],
            ["X2", "Delisted by the exchange after a bankruptcy filing", "as X1; blank when there was no OTC trade "
                                                                       "within 10 trading days"]],
           [6, 48, 70])
    ws3 = wb.create_sheet("About")
    _sheet(ws3, ["item", "value"], [["Product", "Merriden End-of-Day US Equity: delisting values"],
                                     ["Coverage", "Securities in list PBK-RU-01 delisted 2 January 2018 to 30 "
                                                  "September 2026"],
                                     ["Prepared", "7 October 2026"]], [14, 80], freeze=None)
    wb.save(path)


def write_recons(path, recons, changes):
    wb = Workbook()
    first = True
    for rc in recons:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = rc["label"]
        head = ["Constituent", "Ticker", "CUSIP", "Sector", "Close (USD)", "Shares in Index", "Float Factor",
                "Capping Factor", "Index Market Value (USD)", "Weight (%)"]
        data = [[r["name"], r["ticker"], r["cusip"], r["sector"], r["close"], r["shares_in_index"], r["float_factor"],
                 r["capping_factor"], r["mcap"], r["weight"]] for r in rc["rows"]]
        _sheet(ws, head, data, [40, 8, 11, 22, 11, 15, 11, 12, 20, 10],
               {"E": "#,##0.00", "F": "#,##0", "G": "0.000", "H": "0.000000", "I": "#,##0", "J": "0.0000"})
    ws = wb.create_sheet("Interim Changes")
    _sheet(ws, ["Effective (after close)", "Action", "Constituent", "Ticker", "CUSIP", "Reason", "Shares in Index"],
           [[c["date"], c["action"], c["name"], c["ticker"], c["cusip"], c["reason"], c.get("shares")]
            for c in changes], [20, 9, 40, 8, 11, 26, 15], {"A": "yyyy-mm-dd", "G": "#,##0"})
    ws = wb.create_sheet("Notes")
    _sheet(ws, ["Note"], [["Jun2017 lists the constituents in effect on 2 January 2018 under the June 2017 "
                           "reconstitution."],
                          ["Each later sheet lists the constituents at the close of the reconstitution date, the "
                           "fourth Friday of June."],
                          ["Weights are at that close and drift with prices until the next reconstitution."]],
           [110], freeze=None)
    wb.save(path)


# ------------------------------------------------------------------------------------ the README
def write_readme(path, counts):
    f = W.FILES
    lines = [
        "# Research store extracts, US small and mid cap equities",
        "",
        "Systematic Equity, research data. Pulled for the store refresh of 9 October 2026.",
        "",
        "| File | Source | Pulled | Rows |",
        "|---|---|---|---|",
    ]
    src = {"px_old": "Merriden, legacy platform (final history file)", "px_new": "Merriden, current platform",
           "secmaster": "Merriden", "ca": "Merriden", "delist": "Merriden", "fund": "Ledgerpoint",
           "earn": "Ledgerpoint", "recon": "Northcote", "levels": "Northcote", "si": "Exchange short interest files",
           "sweep": "Harrowgate Trust (custodian)", "orders": "Pellbrook OMS", "fills": "Kelloway Securities",
           "dict_px": "Merriden", "spec_ca": "Merriden", "guide_fund": "Ledgerpoint", "method_idx": "Northcote",
           "fees": "Kelloway Securities"}
    for k, name in f.items():
        if k == "readme":
            continue
        lines.append("| %s | %s | %s | %s |" % (name, src[k], W.AS_OF[k], "{:,}".format(counts[k])
                                                if counts.get(k) is not None else ""))
    lines += [
        "",
        "## Notes",
        "",
        "- Universe: security list PBK-RU-01, every US common stock the store has followed since January 2018, "
        "delisted names included.",
        "- Prices: the legacy history file ends 29 Dec 2023. The current platform's Parquet starts 2 Oct 2023, so Q4 "
        "2023 is in both. Keyed on security_id / SID (same number).",
        "- Short interest is the exchanges' published file stacked as received, by symbol, with revisions as "
        "published.",
        "- OMS times are UTC. Kelloway's executions are New York time.",
        "- Sweep rate is the custodian's posted rate for each calendar day (weekends and holidays carry the previous "
        "business day's rate). Reference rate changed 1 Jul 2023.",
        "- Ledgerpoint keys on lp_id; there is no Merriden id in its files. Map through ticker or CUSIP as of the "
        "date.",
        "",
    ]
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        fh.write("\n".join(lines))
