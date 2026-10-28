"""task174 documents: Pellmore's skill-builder guide and national benchmark note (pdf), the district's spring plan
(docx), the Math Lab procedures (pdf) and the budget office's message on the reteach places (eml)."""
from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P

PL = "{:,}".format(P.PLACES)


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.8, leading=13.6, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=14, leading=18,
                        spaceAfter=3, textColor=colors.HexColor("#22405F"))
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.6, leading=14,
                        spaceBefore=8, spaceAfter=4, keepWithNext=1)
    small = ParagraphStyle("small", parent=body, fontSize=8.2, leading=10.8, textColor=colors.HexColor("#444444"))
    return body, h1, h2, small


def _footer(text):
    def foot(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.4)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(18 * mm, 10 * mm, text)
        canvas.drawRightString(LETTER[0] - 18 * mm, 10 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return foot


def _table(rows, widths, size=8.6):
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=size, leading=size + 2.4)
    head = ParagraphStyle("head", parent=cell, fontName="Helvetica-Bold")
    rows = [[Paragraph(str(v), head if r == 0 else cell) for v in row] for r, row in enumerate(rows)]
    t = Table(rows, colWidths=[w * mm for w in widths], repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9A9A9A")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E7ECF2")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


def _doc(path, title, author, flow, footer, when):
    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title=title, author=author, creator=author, invariant=1)
    doc.build(flow, onFirstPage=_footer(footer), onLaterPages=_footer(footer))
    b = open(path, "rb").read()
    fixed = b"D:20000101000000+00'00'"
    assert len(when) == len(fixed) and b.count(fixed) >= 1, (path, b.count(fixed))
    open(path, "wb").write(b.replace(fixed, when.encode()))


def guide(path):
    body, h1, h2, small = _styles()
    s = [Paragraph("Skill builders: how they work and what the district export holds", h1),
         Paragraph("%s help centre, article RM-214. Last updated 11 August 2025." % P.PRODUCT, small),
         Spacer(1, 4),
         Paragraph("What a skill builder is", h2),
         Paragraph("A skill builder gives a student problems on one skill. Each problem is drawn at random from the "
                   "skill builder's pool of problems written to the same standard and the same difficulty, so two "
                   "students rarely see the same problem in the same order. A student completes the skill builder by "
                   "answering three problems in a row correctly on the first attempt. Once it is complete the "
                   "student is not served any more of its problems.", body),
         Paragraph("A problem counts as right only when the student's first action is a correct answer. Asking for a "
                   "hint first, or answering wrong first, counts the problem as wrong. After a wrong answer the "
                   "student may be shown scaffold questions that break the problem into steps, and then moves on to "
                   "the next problem. Scaffold questions are not problems: they do not count toward three in a row "
                   "or toward the daily limit.", body),
         Paragraph("Homework mode", h2),
         Paragraph("Homework mode is the default. The teacher sets the date the skill builder opens and the date it "
                   "closes, and the student can start at most ten problems of the skill builder in one day. The "
                   "count resets at midnight in the school's time zone. After the close date the skill builder stops "
                   "serving problems, complete or not.", body),
         Paragraph("Class mode", h2),
         Paragraph("In class mode there is no daily limit. The teacher sets a fixed number of problems, which every "
                   "student answers whatever their run of right answers, and can add check problems from the same "
                   "skill builder that the student answers straight after the set. The platform reports the number "
                   "of check problems each student got right.", body),
         Paragraph("The district problem export", h2),
         Paragraph("The district export has one row for every problem a student starts: its assignment, the "
                   "student's first action, whether that first action was a correct answer, the hints and attempts "
                   "used, the times the student started and finished it, and the student's school and section. "
                   "Scaffold questions are not in the export. Exports from the legacy reporting tool, used until 31 "
                   "December 2025, give times in the school's local time. Exports from the Pellmore data platform, "
                   "which holds records from 1 January 2026, give times in UTC and use the platform's own column "
                   "names.", body),
         Spacer(1, 4),
         Paragraph("Student accounts are created and moved between schools by the district's nightly roster sync. "
                   "A student who leaves the district has the account deactivated; a student who later enrolls again "
                   "is given a new account.", body)]
    _doc(path, "Skill builders: how they work and what the district export holds", P.VENDOR, s,
         "%s. Help article RM-214." % P.VENDOR, "D:20250811102007-04'00'")


def benchmarks(path, rows):
    body, h1, h2, small = _styles()
    s = [Paragraph("National benchmarks 2025-26: grade 7 homework skill builders", h1),
         Paragraph("%s research team. Published 30 July 2026. All partner districts, homework mode only." % P.VENDOR, small),
         Spacer(1, 4),
         Paragraph("The table gives, for each grade 7 skill builder in the core sequence, the share of assigned "
                   "students who completed it by the close date and the median minutes completing students spent "
                   "on it, across partner districts in the 2025-26 school year. Districts use it to compare their "
                   "own completion with the partner average.", body),
         _table([["Skill builder", "Title", "Students assigned", "Completed by close", "Median minutes to complete"]]
                + rows, [22, 62, 26, 26, 30]),
         Spacer(1, 4),
         Paragraph("Partner districts set their own open and close dates; most leave a skill builder open for two "
                   "school weeks.", small)]
    _doc(path, "National benchmarks 2025-26: grade 7 homework skill builders", P.VENDOR, s,
         "%s. National benchmarks 2025-26." % P.VENDOR, "D:20260730155912-04'00'")


def lab_procedures(path):
    body, h1, h2, small = _styles()
    name, title = P.PEOPLE["interventions"]
    s = [Paragraph("Math Lab procedures, 2025-26", h1),
         Paragraph("%s, Interventions Program. Issued 18 August 2025." % P.DISTRICT, small),
         Spacer(1, 4),
         Paragraph("Who comes to Math Lab", h2),
         Paragraph("Math Lab runs in each middle school's WIN period on the ten school days after a homework skill "
                   "builder closes. Students who had not completed the skill builder by its close are expected to "
                   "attend until they have worked through it with the interventionist. Sixth, seventh and eighth "
                   "graders share the room.", body),
         Paragraph("Recording a session", h2),
         Paragraph("The interventionist records each session in the intervention tool on the day: the date, the "
                   "period, the length of the session in minutes and every student who attended, with the grade and "
                   "unit each student worked on. The tool exports one row per student per session save. When a "
                   "session record is edited after it is first saved (a student added, the length corrected), the "
                   "tool saves the session again and keeps every save; the latest save is the session's record.", body),
         Paragraph("Hours", h2),
         Paragraph("Interventionist hours are reported to the Interventions Program each month from the session "
                   "records, a session's minutes counting once for each interventionist who led it.", body)]
    _doc(path, "Math Lab procedures, 2025-26", P.DISTRICT, s, "%s. Interventions Program." % P.DISTRICT,
         "D:20250818083641-04'00'")


def plan(path):
    name, title = P.PEOPLE["cao"]
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    d.add_heading("Seventh-grade math practice in the spring semester", level=1)
    d.add_paragraph("To: middle school principals and mathematics department chairs\n"
                    "From: %s, %s\nCopy: %s; %s\nDate: 21 October 2026"
                    % (name, title, P.PEOPLE["coordinator"][0], P.PEOPLE["interventions"][0]))
    paras = [
        "From Tuesday 5 January 2027, seventh-grade math will do the Pellmore skill builders for Units 4 to 7 in class "
        "instead of as homework. Each unit has one skill builder, and each skill builder takes one class period in "
        "class mode: every student answers a set of problems from the skill builder and then five more problems from "
        "the same skill builder as the check, in the same period. A student who gets fewer than three of the five "
        "check problems right joins the next school day's reteach group with an interventionist. A student absent "
        "that day does the set and the check in the period after they return.",
        "The board's budget amendment funds %s reteach places for Units 4 to 7 together. A place is one student "
        "in one unit's reteach group, so a student who falls short on two units' checks uses two places." % PL,
        "The set length is the math office's decision and goes into the spring pacing guide we print on 4 December. "
        "The period has room for a set of 6, 8 or 10 problems and the check, and every class will use the same "
        "length. We will use the shortest set whose expected failed checks across the four units fit within the "
        "%s places. Reteach staffing is planned on the seventh graders on roll on 23 October 2026." % PL,
        "Nothing changes in the fall. Unit 3's two skill builders stay as homework, open for ten school days, with "
        "Math Lab for students who have not completed them by the close.",
    ]
    for p in paras:
        d.add_paragraph(p)
    d.core_properties.author = name
    d.core_properties.title = "Seventh-grade math practice in the spring semester"
    d.save(path)


def amendment_email():
    gname, gtitle = P.PEOPLE["budget"]
    cname, ctitle = P.PEOPLE["cao"]
    lname, ltitle = P.PEOPLE["coordinator"]
    bname, btitle = P.PEOPLE["interventions"]
    mail = lambda n: n.lower().replace(" ", ".") + "@tamberlin.k12.nc.us"
    return ("From: %s <%s>\r\nTo: %s <%s>\r\nCc: %s <%s>, %s <%s>\r\n"
            "Date: Wed, 14 Oct 2026 08:52:00 -0400\r\nSubject: Spring reteach places: board amendment approved\r\n"
            "Message-ID: <20261014085200.4471@tamberlin.k12.nc.us>\r\nMIME-Version: 1.0\r\n"
            "Content-Type: text/plain; charset=utf-8\r\nContent-Transfer-Encoding: 8bit\r\n\r\n"
            "Tara,\r\n\r\nThe board approved budget amendment 2026-14 at last night's meeting. It funds %s reteach "
            "places for seventh-grade Units 4 to 7 in the spring, at the place rate in the interventionist contract. "
            "That is the ceiling for the semester: there is no contingency line behind it, so the set length the math "
            "office picks has to fit inside it.\r\n\r\nThe contract runs from 5 January to 4 June 2027. I have asked "
            "Claire to keep the monthly hours report going in the same format as Math Lab so we can compare the two "
            "programs at year end.\r\n\r\nGreg\r\n\r\n%s\r\n%s\r\n"
            % (gname, mail(gname), cname, mail(cname), lname, mail(lname), bname, mail(bname), PL, gname, gtitle))
