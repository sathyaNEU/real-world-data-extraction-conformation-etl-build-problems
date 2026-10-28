"""Write the task174 pack into target/ from the simulated world."""
import csv
import datetime as dt
import os
import zoneinfo

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import docs
import params as P
import postprocess
from world import World, d

EAST = zoneinfo.ZoneInfo(P.TZ)
UTC = dt.timezone.utc
SCHOOL = {s[0]: s for s in P.SCHOOLS}
SBS = {b[0]: b for b in P.SKILL_BUILDERS}

FILES = {
    "log_legacy": "pellmore_problem_log_2025-26_aug_dec.csv",
    "log_2026": "pellmore_problem_log_2026_jan_jun.csv",
    "log_fall": "pellmore_problem_log_2026-27_to_2026-10-23.csv",
    "assignments": "pellmore_assignments_2025-2027.csv",
    "accounts": "pellmore_accounts_2026-10-23.csv",
    "link": "pellmore_student_link.csv",
    "roster": "grade7_roster_2026-10-23.xlsx",
    "archive": "sis_archive/enrollment_history_2025-26.csv",
    "pacing": "grade7_math_pacing_2025-2027.xlsx",
    "guide": "pellmore_skill_builder_guide.pdf",
    "plan": "spring_2027_grade7_math_plan.docx",
    "email": "reteach_places_board_amendment.eml",
    "lab": "math_lab_sessions_2025-26.csv",
    "lab_proc": "math_lab_procedures_2025-26.pdf",
    "eog": "state_math_results_grade7_2026_by_school.csv",
    "bench": "pellmore_national_benchmarks_2025-26.pdf",
}
PACK = list(FILES)
REASON = {"T1": "Transferred to another district school", "W2": "Withdrew from the district",
          "EOY": "End of school year"}
DISTRACTORS = ["eog", "bench", "lab", "lab_proc"]


def _csv(path, header, rows):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(header)
        w.writerows(rows)


def _xlsx(path, sheets, made, saved, person):
    wb = Workbook()
    for i, (title, header, rows, widths) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        if header:
            ws.append(header)
            for c in ws[1]:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="DCE3EC")
                c.alignment = Alignment(vertical="top", wrap_text=True)
        for r in rows:
            ws.append(list(r))
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        if header:
            ws.freeze_panes = "A2"
    wb.properties.creator = person
    wb.properties.lastModifiedBy = person
    wb.save(path)
    postprocess.scrub(path, P.DISTRICT, saved[:10])
    postprocess.finish_office(path, made, saved, P.DISTRICT, person=person)


def utc(t):
    return t.replace(tzinfo=EAST).astimezone(UTC)


class Sections:
    """each student's section at each school, fixed by a seeded draw"""
    def __init__(self, w):
        rng = np.random.default_rng(P.SEED + 7)
        self.n = {code: int(np.ceil(n / 26.0)) for code, name, n, off, win in P.SCHOOLS}
        self.of = {}
        for year in P.YEARS:
            for s, spells, accts in w.cohorts[year]:
                for code, a, b, why in spells:
                    self.of[(s["id"], code)] = "M7-%02d" % (1 + int(rng.integers(0, self.n[code])))

    def get(self, sid, code):
        return self.of.get((sid, code), "M7-01")


def build(target):
    os.makedirs(target, exist_ok=True)
    F = {k: os.path.join(target, v) for k, v in FILES.items()}
    w = World()
    sec = Sections(w)
    rng = np.random.default_rng(P.SEED + 11)

    # assignments: one per section and skill builder in each year (homework mode)
    legacy_id, ref = {}, {}
    rows_a = []
    next_id = 7_340_118
    ids = np.random.default_rng(P.SEED + 13)
    sec_id = {}
    used = set()
    for year in P.YEARS:
        for code, name, n, off, win in P.SCHOOLS:
            for k in range(1, sec.n[code] + 1):
                v = int(ids.integers(10000, 99999))
                while v in used:
                    v = int(ids.integers(10000, 99999))
                used.add(v)
                sec_id[(year, code, "M7-%02d" % k)] = "S%d" % v
    seen_ref = set()
    for year in P.YEARS:
        for sb, (o, c) in sorted(P.WINDOWS[year].items(), key=lambda kv: kv[1][0]):
            for code, name, n, off, win in P.SCHOOLS:
                for k in range(1, sec.n[code] + 1):
                    section = "M7-%02d" % k
                    key = (year, sb, code, section)
                    next_id += int(rng.integers(3, 60))
                    legacy_id[key] = next_id
                    r_ = "AS%08X" % int(ids.integers(0, 2 ** 32))
                    while r_ in seen_ref:
                        r_ = "AS%08X" % int(ids.integers(0, 2 ** 32))
                    seen_ref.add(r_)
                    ref[key] = r_
                    rows_a.append([next_id if year == "2025-26" else "", ref[key], year, sb, SBS[sb][1], SBS[sb][2],
                                   "homework", sec_id[(year, code, section)], o, c])
    _csv(F["assignments"], ["legacy_assignment_id", "assignment_ref", "school_year", "skill_builder", "unit",
                            "skill_builder_title", "mode", "section_id", "opens", "closes"], rows_a)

    # problem logs
    bases = np.random.default_rng(P.SEED + 17)
    pool_base, at = {}, 3170
    for sb in sorted(SBS):
        pool_base[sb] = at
        at += 540 + int(bases.integers(0, 900))
    legacy, new26, fall = [], [], []
    ex_legacy, ex_2026, ex_fall = d(P.EXPORT_DAY["legacy"]), d(P.EXPORT_DAY["2025-26"]), d(P.EXPORT_DAY["2026-27"])
    cut = dt.datetime.fromisoformat(P.LEGACY_UNTIL)
    for e in sorted(w.events, key=lambda e: (e["start"], e["user_id"])):
        sb = e["unit"]
        if e["student_id"] is None:
            section = "M7-01"
        else:
            section = sec.get(e["student_id"], e["school"])
        key = (e["year"], sb, e["school"], section)
        if e["original"]:
            pnum, pref = pool_base[sb] + e["pid"], "PRB-%s-%04d" % (sb, e["pid"])
        else:
            parent = e.get("pid_parent", 0)
            pnum, pref = pool_base[sb] + 400 + (-e["pid"]), "PRB-%s-S%d" % (sb, -e["pid"])
        first = "hint" if e["hint"] else "answer"
        if not e["original"]:
            continue
        if e["year"] == "2025-26" and e["start"] < cut:
            school = SCHOOL[w.export_school(e["user_id"], ex_legacy)][1]
            legacy.append([e["user_id"], legacy_id[key], pnum,
                           "hint" if e["hint"] else "attempt", int(e["correct"]), (1 if e["hint"] else 0),
                           e["attempts"], e["start"].strftime("%Y-%m-%d %H:%M:%S"), e["end"].strftime("%Y-%m-%d %H:%M:%S"),
                           school, "Math 7 %s" % section.replace("M7-", "P")])
        elif e["original"]:
            day = ex_2026 if e["year"] == "2025-26" else ex_fall
            row = [e["user_id"], ref[key], pref, first, int(e["correct"]),
                   (1 if e["hint"] else 0), e["attempts"], utc(e["start"]).strftime("%Y-%m-%dT%H:%M:%SZ"),
                   utc(e["end"]).strftime("%Y-%m-%dT%H:%M:%SZ"), w.export_school(e["user_id"], day),
                   sec_id.get((e["year"], e["school"], section), sec_id[(e["year"], P.SCHOOLS[0][0], "M7-01")])]
            (new26 if e["year"] == "2025-26" else fall).append(row)
    _csv(F["log_legacy"], ["user_id", "assignment_id", "problem_id", "first_response", "correct",
                           "hint_count", "attempt_count", "start_time", "end_time", "school", "class_name"], legacy)
    head2 = ["student_user_id", "assignment_ref", "problem_ref", "first_action", "first_attempt_score",
             "hints", "attempts", "started_utc", "finished_utc", "school_code", "section_id"]
    _csv(F["log_2026"], head2, new26)
    _csv(F["log_fall"], head2, fall)

    # accounts as of the as-of date, and the district link
    acc_rows, link_rows = [], []
    for a in sorted(w.accounts, key=lambda a: a["user_id"]):
        rec = w.by_user[a["user_id"]]
        if rec is None:
            acc_rows.append([a["user_id"], "student", 7, a["school"], "active", a["opened"].isoformat(), ""])
            continue
        s, spells, accts = rec
        grade = 7 if s["year"] == "2026-27" else 8
        school = w.export_school(a["user_id"], P.AS_OF)
        status = "active" if a["closed"] is None else "deactivated"
        acc_rows.append([a["user_id"], "student", grade, school, status, a["opened"].isoformat(),
                         a["closed"].isoformat() if a["closed"] else ""])
        link_rows.append([s["id"], a["user_id"], a["opened"].isoformat()])
    _csv(F["accounts"], ["user_id", "account_type", "grade", "school_code", "status", "created_on", "deactivated_on"],
         acc_rows)
    _csv(F["link"], ["student_id", "pellmore_user_id", "linked_on"], sorted(link_rows))

    # the SIS: the roll on the as-of date, and the archived 2025-26 enrollment history
    on_roll, left = [], []
    for s, spells, accts in w.cohorts["2026-27"]:
        code = w.school_on(spells, P.AS_OF)
        if code is not None:
            entry = [sp for sp in spells if sp[0] == code and sp[1] <= P.AS_OF][-1][1]
            on_roll.append([s["id"], code, SCHOOL[code][1], 7, entry,
                            "Y" if rng.random() < 0.09 else "N", "Y" if rng.random() < 0.13 else "N"])
        for code_, a_, b_, why in spells:
            if b_ is not None and b_ <= P.AS_OF:
                left.append([s["id"], code_, b_, REASON[why]])
    _xlsx(F["roster"], [("On roll 23 Oct 2026", ["student_id", "school_code", "school", "grade", "first_day_on_roll",
                                                   "el", "iep"], sorted(on_roll), [12, 11, 34, 7, 16, 6, 6]),
                        ("Left since 24 Aug 2026", ["student_id", "school_code", "last_day_on_roll", "reason"],
                         sorted(left), [12, 11, 16, 34])],
          "2026-10-23T13:02:00Z", "2026-10-23T13:40:00Z", P.PEOPLE["analyst"][0])
    arch = []
    end25 = d(P.YEAR_END["2025-26"])
    for s, spells, accts in w.cohorts["2025-26"]:
        for code, a, b, why in spells:
            arch.append([s["id"], "2025-26", code, 7, a.isoformat(), (b or end25).isoformat(), REASON[why or "EOY"]])
    _csv(F["archive"], ["student_id", "school_year", "school_code", "grade", "first_day_on_roll", "last_day_on_roll",
                        "exit_reason"], sorted(arch))

    # pacing
    pace = []
    for year in P.YEARS:
        rows = []
        for sb, (o, c) in sorted(P.WINDOWS[year].items(), key=lambda kv: kv[1][0]):
            rows.append([SBS[sb][1], P.UNIT_TITLES[SBS[sb][1]], sb, SBS[sb][2], "homework", d(o), d(c)])
        if year == "2026-27":
            rows.append([3, P.UNIT_TITLES[3], "3A", SBS["3A"][2], "homework", d("2026-11-02"), d("2026-11-13")])
            rows.append([3, P.UNIT_TITLES[3], "3B", SBS["3B"][2], "homework", d("2026-11-16"), d("2026-12-04")])
            for sb in P.SPRING:
                u0, u1 = P.SPRING_2027[int(sb)]
                rows.append([int(sb), P.UNIT_TITLES[int(sb)], sb, SBS[sb][2], "class (spring plan)", d(u0), d(u1)])
        pace.append((year, ["unit", "unit_title", "skill_builder", "skill_builder_title", "mode", "opens or unit starts",
                            "closes or unit ends"], rows, [6, 26, 12, 40, 18, 18, 18]))
    _xlsx(F["pacing"], pace, "2025-07-28T14:10:00Z", "2026-10-02T09:15:00Z", P.PEOPLE["coordinator"][0])

    # documents
    docs.guide(F["guide"])
    postprocess.scrub(F["guide"], "Microsoft Word 2019", "2025-08-11T10:20:00")
    docs.plan(F["plan"])
    postprocess.scrub(F["plan"], P.DISTRICT, "2026-10-21")
    postprocess.finish_office(F["plan"], "2026-10-20T15:12:00Z", "2026-10-21T09:48:00Z", P.DISTRICT,
                              person=P.PEOPLE["cao"][0])
    with open(F["email"], "w", encoding="utf-8", newline="") as fh:
        fh.write(docs.amendment_email())
    docs.lab_procedures(F["lab_proc"])
    postprocess.scrub(F["lab_proc"], "Microsoft Word 2019", "2025-08-18T08:30:00")

    # Math Lab sessions export
    lab_rows = []
    for r in sorted(w.lab, key=lambda r: (r["school"], r["date"], r["saved_at"], r["staff"], r["student_id"])):
        g, u = r["unit"].split("-")
        lab_rows.append([r["school"], r["date"].isoformat(), r["period"], r["staff"], "Grade %s Unit %s" % (g, u),
                         r["student_id"], r["grade"], r["minutes"], r["saved_at"].strftime("%Y-%m-%d %H:%M")])
    _csv(F["lab"], ["school_code", "session_date", "period", "interventionist", "unit_worked", "student_id", "grade",
                    "minutes", "saved_at"], lab_rows)

    # distractors: the state results by school, the vendor's national benchmarks
    eog = []
    for code, name, n, off, win in P.SCHOOLS:
        tested = int(n * 0.97 + rng.integers(-4, 5))
        mean = 1.0 / (1.0 + np.exp(-(0.35 + 0.9 * off)))
        eog.append([code, name, tested, "%.1f" % (448 + 9 * off + rng.normal(0, 0.8)), "%.1f" % (100 * mean),
                    "%.1f" % (100 * mean * 0.62)])
    _csv(F["eog"], ["school_code", "school", "students_tested", "mean_scale_score", "pct_level_3_and_above",
                    "pct_level_4_and_above"], eog)
    bench = []
    for sb in sorted(SBS, key=lambda x: (len(x), x)):
        bench.append([sb, SBS[sb][2], "{:,}".format(int(rng.integers(38000, 52000))),
                      "%d%%" % int(round(82 - 10 * max(0, -SBS[sb][3] - 0.6) + rng.normal(0, 1.5))),
                      int(round(14 + 9 * max(0, -SBS[sb][3]) - 20 * SBS[sb][4] + rng.normal(0, 0.8)))])
    docs.benchmarks(F["bench"], bench)
    postprocess.scrub(F["bench"], "Microsoft Word 2019", "2026-07-30T16:00:00")
    return F, w


if __name__ == "__main__":
    import sys
    import time
    t0 = time.time()
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "target")
    F, w = build(out)
    for k, f in F.items():
        print("%-12s %9d bytes  %s" % (k, os.path.getsize(f), os.path.relpath(f, out)))
    print("%.1fs" % (time.time() - t0))
