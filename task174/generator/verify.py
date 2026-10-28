"""Independent verifier: reads only target.zip (csv, openpyxl, python-docx, pypdf; none of world.py, pack.py, solve.py
or params.py) and recomputes every graded figure on its own code path, then compares with figures.json."""
import csv
import datetime as dt
import io
import json
import os
import re
import sys
import zipfile
import zoneinfo

import numpy as np
import openpyxl
from docx import Document
from pypdf import PdfReader

HERE = os.path.dirname(os.path.abspath(__file__))
EAST = zoneinfo.ZoneInfo("America/New_York")


def main(zpath):
    z = zipfile.ZipFile(zpath)
    names = z.namelist()
    fig = json.load(open(os.path.join(HERE, "figures.json")))
    n = [0]

    def check(label, cond, detail=""):
        n[0] += 1
        if not cond:
            raise AssertionError("VERIFY FAIL %s %s" % (label, detail))

    pick = lambda pre: [x for x in names if x.startswith(pre)][0]

    def rows(name):
        with z.open(name) as f:
            return list(csv.DictReader(io.TextIOWrapper(f, encoding="utf-8")))

    def pdf(name):
        return re.sub(r"\s+", " ", " ".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(z.read(name))).pages))

    # the documents ------------------------------------------------------------------------------------------------
    plan = re.sub(r"\s+", " ", " ".join(p.text for p in Document(io.BytesIO(z.read(pick("spring_2027_grade7_math_plan")))).paragraphs))
    check("docx plan read back", "Units 4 to 7" in plan and len(plan) > 1200)
    places = int(re.search(r"funds ([\d,]+) reteach places", plan).group(1).replace(",", ""))
    lengths = [int(x) for x in re.search(r"a set of (\d+), (\d+) or (\d+) problems", plan).groups()]
    check_n = 5 if "five more problems from the same skill builder as the check" in plan else None
    bar = 3 if "fewer than three of the five check problems right" in plan else None
    check("plan: places, lengths, check and bar parsed", places == fig["places"] and lengths == fig["lengths"]
          and check_n == 5 and bar == 3, (places, lengths))
    check("plan: the roll date", "on roll on 23 October 2026" in plan)
    check("plan: the shortest set fitting the places", "We will use the shortest set whose expected failed checks" in plan)
    guide = pdf(pick("pellmore_skill_builder_guide"))
    check("pdf guide read back", "three problems in a row correctly on the first attempt" in guide)
    check("pdf guide: right means a correct first action", "counts as right only when the student's first action is a "
                                                           "correct answer" in guide)
    check("pdf guide: export times", "give times in UTC" in guide and "local time" in guide)
    eml = z.read(pick("reteach_places_board_amendment")).decode("utf-8")
    check("eml: places agree with the plan", "{:,}".format(places) in eml)

    # the frame: first attempts by student, skill builder and time, sittings by Eastern day ------------------------
    link = {r["pellmore_user_id"]: r["student_id"] for r in rows(pick("pellmore_student_link"))}
    asg = rows(pick("pellmore_assignments"))
    by_legacy = {r["legacy_assignment_id"]: r["skill_builder"] for r in asg if r["legacy_assignment_id"]}
    by_ref = {r["assignment_ref"]: r["skill_builder"] for r in asg}
    recs = []
    for r in rows(pick("pellmore_problem_log_2025-26_aug_dec")):
        if r["user_id"] in link:
            recs.append((link[r["user_id"]], by_legacy[r["assignment_id"]],
                         dt.datetime.strptime(r["start_time"], "%Y-%m-%d %H:%M:%S"), int(r["correct"])))
    for pre in ("pellmore_problem_log_2026_jan_jun", "pellmore_problem_log_2026-27"):
        for r in rows(pick(pre)):
            if r["student_user_id"] in link:
                t = dt.datetime.strptime(r["started_utc"], "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=dt.timezone.utc)
                recs.append((link[r["student_user_id"]], by_ref[r["assignment_ref"]],
                             t.astimezone(EAST).replace(tzinfo=None), int(r["first_attempt_score"])))
    check("frame: 25,000 rows or more", len(recs) >= 25000, len(recs))
    recs.sort(key=lambda x: (x[0], x[1], x[2]))
    sbs = sorted({r[1] for r in recs}, key=lambda s: (len(s), s))
    sbi = {s: i for i, s in enumerate(sbs)}
    stu = {}
    S, B, E1, E2, Y = [], [], [], [], []
    prev, day, e1, e2 = None, None, 0, 0
    for s, sb, t, y in recs:
        if (s, sb) != prev:
            prev, day, e1, e2 = (s, sb), t.date(), 0, 0
        elif t.date() != day:
            e2 += e1
            e1 = 0
            day = t.date()
        S.append(stu.setdefault(s, len(stu)))
        B.append(sbi[sb])
        E1.append(e1)
        E2.append(e2)
        Y.append(y)
        e1 += 1
    S, B, E1, E2, Y = map(np.array, (S, B, E1, E2, Y))
    E1, E2, Y = E1.astype(float), E2.astype(float), Y.astype(float)
    ns, nb = len(stu), len(sbs)

    # block-coordinate Newton: unit parameters by a 3x3 Newton step per skill builder, abilities by a 1-D Newton
    # step per student under a standard normal prior; iterate to a fixed point
    th, U = np.zeros(ns), np.zeros((nb, 3))
    for it in range(400):
        lg = th[S] + U[B, 0] + U[B, 1] * E1 + U[B, 2] * E2
        p = 1.0 / (1.0 + np.exp(-lg))
        r, w = Y - p, p * (1 - p)
        X = [np.ones_like(E1), E1, E2]
        g = np.stack([np.bincount(B, r * x, nb) for x in X], 1)
        H = np.zeros((nb, 3, 3))
        for i in range(3):
            for j in range(3):
                H[:, i, j] = np.bincount(B, w * X[i] * X[j], nb)
        step = np.linalg.solve(H, g[:, :, None])[:, :, 0]
        U = U + step
        lg = th[S] + U[B, 0] + U[B, 1] * E1 + U[B, 2] * E2
        p = 1.0 / (1.0 + np.exp(-lg))
        gs = np.bincount(S, Y - p, ns) - th
        hs = np.bincount(S, p * (1 - p), ns) + 1.0
        dth = gs / hs
        th = th + dth
        if max(np.abs(step).max(), np.abs(dth).max()) < 1e-9:
            break
    check("fit converged", it < 399, it)
    lg = th[S] + U[B, 0] + U[B, 1] * E1 + U[B, 2] * E2
    p = 1.0 / (1.0 + np.exp(-lg))
    info = np.bincount(S, p * (1 - p), ns)

    # the forecast for the seventh graders on roll, integrated over each student's ability ---------------------------
    wb = openpyxl.load_workbook(io.BytesIO(z_read(z, pick("grade7_roster"))), read_only=True)
    roll = [r[0] for r in list(wb["On roll 23 Oct 2026"].iter_rows(values_only=True))[1:] if r and r[0]]
    check("roll: every student once", len(roll) == len(set(roll)) == fig["on_roll"], len(roll))
    idx = np.array([stu.get(s, -1) for s in roll])
    mu = np.where(idx >= 0, th[np.maximum(idx, 0)], 0.0)
    sd = np.where(idx >= 0, 1.0 / np.sqrt(info[np.maximum(idx, 0)] + 1.0), 1.0)
    xk, wk = np.polynomial.hermite_e.hermegauss(41)
    wk = wk / wk.sum()

    def short(zz):
        q = 1.0 / (1.0 + np.exp(-zz))
        dist = np.zeros((zz.shape[0], check_n + 1))
        dist[:, 0] = 1
        for k in range(check_n):
            qk = q[:, k:k + 1]
            nd = dist * (1 - qk)
            nd[:, 1:] += dist[:, :-1] * qk
            dist = nd
        return dist[:, :bar].sum(1)
    spring = [s for s in sbs if s in ("4", "5", "6", "7")]
    tot = {}
    unit8 = {}
    for L in lengths:
        tot[L] = 0.0
        for sb in spring:
            b = sbi[sb]
            base = U[b, 0] + U[b, 1] * np.arange(L, L + check_n)
            acc = sum(wgt * short((mu + sd * x)[:, None] + base[None, :]) for x, wgt in zip(xk, wk))
            tot[L] += acc.sum()
            if L == fig["chosen"]:
                unit8[sb] = acc.sum()
    unit = fig["rounding"]
    r100 = lambda v: int(unit * int((v + unit / 2.0) // unit))
    for L in lengths:
        check("total at %d" % L, r100(tot[L]) == fig["total"][str(L)], (tot[L], fig["total"][str(L)]))
    for sb in spring:
        check("unit %s at %d" % (sb, fig["chosen"]), r100(unit8[sb]) == fig["unit"][sb], (unit8[sb], fig["unit"][sb]))
    chosen = min([L for L in lengths if tot[L] <= places])
    check("the call: the shortest length within the places", chosen == fig["chosen"], chosen)

    # the ask: 2025-26 problems on Units 4 to 7 by the school attended that day ---------------------------------------
    spells = {}
    for r in rows(pick("sis_archive/enrollment_history")):
        spells.setdefault(r["student_id"], []).append((r["school_code"], dt.date.fromisoformat(r["first_day_on_roll"]),
                                                       dt.date.fromisoformat(r["last_day_on_roll"])))
    count = {}
    for r in rows(pick("pellmore_problem_log_2026_jan_jun")):
        sb = by_ref[r["assignment_ref"]]
        if sb not in ("4", "5", "6", "7") or r["student_user_id"] not in link:
            continue
        d0 = dt.datetime.strptime(r["started_utc"], "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=dt.timezone.utc).astimezone(EAST).date()
        at = [c for c, a, b in spells.get(link[r["student_user_id"]], []) if a <= d0 <= b]
        if at:
            count[(at[0], sb)] = count.get((at[0], sb), 0) + 1
    for code, units in fig["problems"].items():
        for sb, v in units.items():
            check("ask %s unit %s" % (code, sb), count.get((code, sb), 0) == v, (count.get((code, sb), 0), v))
    print("verify: %d assertions passed on %s" % (n[0], os.path.basename(zpath)))
    return 0


def z_read(z, name):
    return z.read(name)


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(HERE), "target.zip")))
