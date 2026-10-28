"""The golden deliverables: the set-length workbook, the memo with the case for the length, and the chart."""
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

import params as P  # noqa: E402
import postprocess  # noqa: E402

XLSX = "spring_practice_set_length.xlsx"
DOCX = "spring_reteach_memo.docx"
PNG = "spring_reteach_forecast.png"
AUTHOR = P.PEOPLE["coordinator"][0]
SB_TITLE = {b[0]: b[2] for b in P.SKILL_BUILDERS}


def n(x):
    return "{:,}".format(int(x))


def rr(x):
    return int(250 * int((x + 125) // 250))


WORD = {6: "six", 8: "eight", 10: "ten"}


def _page_field(par):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = par.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def chart_png(path, F):
    lengths = F["lengths"]
    colours = {"4": "#2E5C8A", "5": "#C8553D", "6": "#6A8D73", "7": "#D9A441"}
    fig, ax = plt.subplots(figsize=(7.6, 5.0), dpi=150)
    bottom = [0.0] * len(lengths)
    for sb in P.SPRING:
        vals = [F["unit_all_raw"][str(L)][sb] for L in lengths]
        ax.bar([str(L) for L in lengths], vals, bottom=bottom, color=colours[sb], width=0.58,
               label="Unit %s: %s" % (sb, SB_TITLE[sb]), edgecolor="white", linewidth=0.6)
        bottom = [b + v for b, v in zip(bottom, vals)]
    ax.axhline(F["places"], color="#222222", linestyle="--", linewidth=1.3)
    ax.text(-0.42, F["places"] + 70, "Funded reteach places: %s" % n(F["places"]), ha="left", va="bottom", fontsize=8.6)
    for i, L in enumerate(lengths):
        weight = "bold" if L == F["chosen"] else "normal"
        ax.text(i, bottom[i] + 60, "about %s" % n(F["total"][str(L)]), ha="center", va="bottom", fontsize=9,
                fontweight=weight)
    i = lengths.index(F["chosen"])
    ax.bar([str(F["chosen"])], [bottom[i]], color="none", edgecolor="#111111", linewidth=1.8, width=0.58)
    ax.set_xlabel("Problems in the in-class set before the five-problem check")
    ax.set_ylabel("Expected failed checks, Units 4 to 7")
    ax.set_title("%s problems is the shortest set that keeps expected failed checks inside the %s places"
                 % (WORD[F["chosen"]].capitalize(), n(F["places"])), fontsize=10.5)
    ax.set_ylim(0, max(bottom) * 1.16)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, p: n(v)))
    ax.legend(fontsize=7.8, loc="upper center", bbox_to_anchor=(0.5, -0.15), ncol=2, frameon=False)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    fig.tight_layout()
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)


def xlsx(path, F):
    wb = Workbook()
    ws = wb.active
    ws.title = "Set length"
    head = Font(bold=True)
    fill = PatternFill("solid", fgColor="DCE3EC")
    pick = PatternFill("solid", fgColor="FFF2CC")
    ws.append(["%s-problem sets keep spring 2027 failed checks inside the %s reteach places"
               % (WORD[F["chosen"]].capitalize(), n(F["places"]))])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Seventh graders on roll on 23 October 2026: %s. A failed check is fewer than three of the five check "
               "problems right. Figures to the nearest 250." % n(F["on_roll"])])
    ws.append([])
    ws.append(["Set length (problems)", "Expected failed checks, Units 4 to 7", "Funded places", "Within the places"])
    for c in ws[4]:
        c.font = head
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for L in F["lengths"]:
        ws.append([L, F["total"][str(L)], F["places"], "Yes" if F["total_raw"][str(L)] <= F["places"] else "No"])
        if L == F["chosen"]:
            for c in ws[ws.max_row]:
                c.fill = pick
    ws.append([])
    ws.append(["At %d problems, by unit" % F["chosen"], "Expected failed checks"])
    for c in ws[ws.max_row]:
        c.font = head
        c.fill = fill
    for sb in P.SPRING:
        ws.append(["Unit %s: %s" % (sb, SB_TITLE[sb]), F["unit"][sb]])
    ws.append([])
    ws.append(["Recommended set length", F["chosen"]])
    ws.cell(ws.max_row, 1).font = head
    ws.freeze_panes = "B5"
    for j, w in enumerate([44, 22, 14, 16], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for r in ws.iter_rows(min_row=5):
        for c in r[1:3]:
            if isinstance(c.value, int):
                c.number_format = "#,##0"

    ws2 = wb.create_sheet("Practice 2025-26 Units 4-7")
    ws2.append(["Skill-builder problems seventh graders started on Units 4 to 7 in 2025-26, by the school the student "
                "attended on the day"])
    ws2["A1"].font = Font(bold=True)
    ws2.append([])
    ws2.append(["School code", "School"] + ["Unit %s" % sb for sb in P.SPRING])
    for c in ws2[3]:
        c.font = head
        c.fill = fill
    for code in sorted(F["problems"]):
        ws2.append([code, F["school_names"][code]] + [F["problems"][code][sb] for sb in P.SPRING])
    ws2.append(["", "All schools"] + [sum(F["problems"][c][sb] for c in F["problems"]) for sb in P.SPRING])
    for c in ws2[ws2.max_row]:
        c.font = head
    ws2.append([])
    ws2.append(["", "All four units, all schools", F["problems_total"]])
    ws2.freeze_panes = "C4"
    for r in ws2.iter_rows(min_row=4):
        for c in r[2:]:
            c.number_format = "#,##0"
    for j, w in enumerate([12, 34, 11, 11, 11, 11], 1):
        ws2.column_dimensions[get_column_letter(j)].width = w

    ws3 = wb.create_sheet("Notes")
    notes = [
        ["Basis", "The %s seventh graders on roll on 23 October 2026, the roll the spring plan staffs on." % n(F["on_roll"])],
        ["Failed check", "Fewer than three of the five check problems right, a right answer being a correct first "
                         "action with no hint (Pellmore help article RM-214)."],
        ["Places", "One place is one student in one unit's reteach group; a student short on two units uses two."],
        ["Forecast", "Each student's chance of a failed check after a set done in one class period, from first "
                     "attempts in the 2025-26 and fall 2026 homework logs, with the problems from earlier evenings "
                     "counted apart from the current sitting's; summed over the roll and rounded once to the "
                     "nearest 250."],
        ["Practice table", "Problems started on Units 4 to 7 skill builders between January and June 2026, at the "
                           "school the student attended that day (SIS enrollment history, Eastern time). Staff test "
                           "accounts and problems started between two enrollments are in no school."],
        ["Prepared", "%s, Mathematics Office, 27 October 2026." % AUTHOR],
    ]
    ws3.append(["Item", "Note"])
    for c in ws3[1]:
        c.font = head
        c.fill = fill
    for r in notes:
        ws3.append(r)
    for r in ws3.iter_rows(min_row=2):
        r[1].alignment = Alignment(wrap_text=True, vertical="top")
    ws3.column_dimensions["A"].width = 16
    ws3.column_dimensions["B"].width = 96
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.save(path)
    postprocess.scrub(path, P.DISTRICT, "2026-10-27")
    postprocess.finish_office(path, "2026-10-26T13:20:00Z", "2026-10-27T16:45:00Z", P.DISTRICT, person=AUTHOR)


def docx(path, F, png):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2.2)
    t = d.add_paragraph()
    r = t.add_run("Eight-problem sets keep the spring reteach groups inside the funded places"
                  if F["chosen"] == 8 else "Spring 2027 in-class skill builders: the set length")
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor(0x22, 0x40, 0x5F)
    d.add_paragraph("To: %s, %s\nFrom: %s, %s\nDate: 27 October 2026\nSubject: Set length for the spring in-class "
                    "skill builders (Units 4 to 7)" % (P.PEOPLE["cao"][0], P.PEOPLE["cao"][1], AUTHOR,
                                                       P.PEOPLE["coordinator"][1]))
    L, T = F["chosen"], F["total"]
    p = d.add_paragraph()
    p.add_run("Recommendation. ").bold = True
    p.add_run("Set the spring skill builders at %d problems before the five-problem check. We expect about %s failed "
              "checks across Units 4 to 7, inside the %s reteach places the amendment funds. Six problems would "
              "produce about %s, more than the places; ten would bring the figure down to about %s and spend two "
              "more problems of every period for a cut we do not need."
              % (L, n(T[str(L)]), n(F["places"]), n(T["6"]), n(T["10"])))
    d.add_paragraph().add_run("Why last year's homework cannot simply be carried over").bold = True
    d.add_paragraph(
        "The spring check comes at the end of one class period, straight after the set. Nothing in the homework "
        "record looks like that. At home a student stops the moment they have three right in a row, so the students "
        "still answering a skill builder's ninth or tenth problem are the ones who were struggling with it; the "
        "accuracy of those late problems describes them, not the class. And the students who needed more than one "
        "evening for a skill builder came back to it well below where they had finished the evening before. Read as one learning curve, that drop looks like slow learning. It is not what a "
        "student will do inside one period, where there is no night in between.")
    d.add_paragraph(
        "So the forecast separates the two. Using every seventh grader's first attempts in 2025-26 and this fall, it "
        "estimates each student's level and, for each skill builder, how much a student gains with each problem "
        "inside a sitting, counting the problems from earlier evenings apart. It then works out, for each of the %s "
        "seventh graders on roll on 23 October, the chance of fewer than three of the five check problems right "
        "after a set of 6, 8 or 10 problems in one sitting, allowing for how much we know about each student from "
        "this fall, and adds them up." % n(F["on_roll"]))
    d.add_paragraph(
        "The two shortcuts would have sent us the wrong way. Replaying last year's homework, where most students "
        "eventually finished, makes even six problems look comfortably inside the places. A learning curve "
        "fitted across evenings reads the overnight loss as slow learning and puts every length over the places.")
    d.add_paragraph().add_run("Expected failed checks (nearest 250)").bold = True
    tb = d.add_table(rows=1, cols=3)
    tb.style = "Light Grid Accent 1"
    for i, h in enumerate(["Set length", "Expected failed checks, Units 4 to 7", "Within the 4,100 places"]):
        tb.rows[0].cells[i].text = h
    for Lx in F["lengths"]:
        cells = tb.add_row().cells
        cells[0].text = "%d problems%s" % (Lx, " (recommended)" if Lx == L else "")
        cells[1].text = n(T[str(Lx)])
        cells[2].text = "Yes" if F["total_raw"][str(Lx)] <= F["places"] else "No"
    d.add_paragraph()
    tu = d.add_table(rows=1, cols=2)
    tu.style = "Light Grid Accent 1"
    tu.rows[0].cells[0].text = "At %d problems" % L
    tu.rows[0].cells[1].text = "Expected failed checks"
    for sb in P.SPRING:
        cells = tu.add_row().cells
        cells[0].text = "Unit %s: %s" % (sb, SB_TITLE[sb])
        cells[1].text = n(F["unit"][sb])
    src = d.add_paragraph("Source: Pellmore problem logs, August 2025 to 23 October 2026; SIS roll of 23 October 2026. "
                          "Figures rounded to the nearest 250; a place is one student in one unit's reteach group, so "
                          "the four units add.")
    src.runs[0].font.size = Pt(8.5)
    src.runs[0].italic = True
    d.add_paragraph(
        "At %d problems Unit 5, %s, carries the most, about %s: students gain little on it with each extra problem, "
        "so a longer set helps it least. Unit 4, %s, starts hardest but students pick it up quickly within a "
        "sitting, and by %d problems it is down to about %s." % (L, SB_TITLE["5"], n(F["unit"]["5"]), SB_TITLE["4"], L,
                                                               n(F["unit"]["4"])))
    d.add_picture(png, width=Cm(15.5))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_paragraph().add_run("Context in the workbook").bold = True
    d.add_paragraph(
        "The workbook also lists each school's 2025-26 skill-builder problems on each of Units 4 to 7, counted at the "
        "school each student attended on the day the problem was started (%s in all). It does not change the length; "
        "it shows principals how much homework practice the in-class sets replace." % n(F["problems_total"]))
    d.add_paragraph().add_run("What could move it").bold = True
    d.add_paragraph(
        "The forecast assumes students work through a set in class the way they worked through one evening at home. "
        "Unit 4's checks, from 5 January, are the first test of that.")
    sec = d.sections[0]
    fp = sec.footer.paragraphs[0]
    fp.text = "Tamberlin County Schools, Mathematics Office. Page "
    _page_field(fp)
    for rn in fp.runs:
        rn.font.size = Pt(8)
    d.core_properties.author = AUTHOR
    d.core_properties.title = "Spring 2027 in-class skill builders: the set length"
    d.save(path)
    postprocess.scrub(path, P.DISTRICT, "2026-10-27")
    postprocess.finish_office(path, "2026-10-26T15:02:00Z", "2026-10-27T17:10:00Z", P.DISTRICT, person=AUTHOR)


def write(golden_dir, F):
    os.makedirs(golden_dir, exist_ok=True)
    png = os.path.join(golden_dir, PNG)
    chart_png(png, F)
    xlsx(os.path.join(golden_dir, XLSX), F)
    docx(os.path.join(golden_dir, DOCX), F, png)
