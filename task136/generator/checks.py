"""task136 generator checks (round 3, the rule 3 rework): the call and its bins; what each past bulk-buy added, read
against suburbs whose campaign had not yet run, with every defensible construction of it and a one-sigma move landing
the same figures; the ladder and every rival route (installed systems per home letterboxed, per house without PV, per
owner-occupied house, per first application, an average campaign, each suburb's own before and after, every other
suburb as the comparison, the hazard of the houses left, a two-way fixed-effects regression, fitted models on the
obvious covariates, assumed shares of additional households); the three-rules block on the clean truth (R1 to R3c);
every semantic fork row (SF1 to SF29); the closure axes; the four asks' stops, hazards and devices; the distractors,
the input gates and the realism sweep. Run by build.py on the built pack; raises on the first failure and prints the
count of assertions passed."""
from __future__ import annotations

import collections
import datetime as dt
import json
import os
import re
import shutil
import tempfile

import numpy as np
import pandas as pd

import calc as C
import common as K

N = [0]
SF_SEEN = {}
GOLD = dict(winner="YRB27", runner="YRB24", excess=1.9, runner_excess=1.7, gap=0.2, bdw=1.7, wtv=0.7)
FIGS = ("excess", "runner_excess", "gap", "bdw", "wtv")


def ok(cond, msg):
    if not cond:
        raise AssertionError(msg)
    N[0] += 1


def r1(x):
    return round(x + 1e-12, 1)


def ten(x):
    return int(round(x / 10.0) * 10)


def bin_room(x, unit):
    """Distance from x to the nearest rounding edge, as a share of the bin (0.5 at the centre)."""
    k = x / unit
    return abs((k - np.floor(k)) - 0.5)


def raw(v):
    w = v["winner"]
    return dict(winner=w, runner=v["runner"], excess=v["excess"][w], runner_excess=v["excess"][v["runner"]],
                gap=v["gap"], bdw=v["add_mw"]["Bindarra Waters"], wtv=v["add_mw"]["Wattleview"])


def graded(v):
    g = raw(v)
    return (g["winner"], g["runner"]) + tuple(r1(g[k]) for k in FIGS)


def golden_tuple():
    return (GOLD["winner"], GOLD["runner"]) + tuple(GOLD[k] for k in FIGS)


def misses_all(v):
    """A route misses the call: it names another feeder, or every graded figure it files sits in another bin."""
    g = graded(v)
    if g[0] != GOLD["winner"]:
        return True
    return all(a != b for a, b in zip(g[2:], golden_tuple()[2:]))


def pdf_text(path):
    from pypdf import PdfReader
    return " ".join(" ".join(pg.extract_text() or "" for pg in PdfReader(path).pages).split())


def with_base(L, base_kw, lift, avgkw, pools, hc_case="hc29"):
    tot = dict(base_kw)
    add = {}
    for name, n in lift.items():
        f = pools[name]["feeder"]
        tot[f] += n * avgkw[f]
        add[name] = n * avgkw[f] / 1000.0
    ex = {f: tot[f] / 1000.0 - L["hca"][f][hc_case] for f in K.FEEDERS}
    order = sorted(K.FEEDERS, key=lambda f: -ex[f])
    return dict(excess=ex, order=order, winner=order[0], runner=order[1], gap=ex[order[0]] - ex[order[1]], lift=lift,
                add_mw=add, fc=tot)


def summed_capacity(reg, date):
    d = pd.Timestamp(date)
    live = reg[reg.status.isin(C.CONNECTED) & (reg.connection_date <= d) & (reg.removal_date.isna() | (reg.removal_date > d))]
    return live.groupby("feeder").kw.sum()


def base_with(reg, cap_fn, trend_kw=None, years=K.YEARS_TO_DESIGN, base_date=K.BASE):
    cap = cap_fn(reg, base_date)
    tr = trend_kw if trend_kw is not None else C.trend(reg)[0]
    return {f: cap.get(f, 0.0) + years * tr[f] for f in K.FEEDERS}


# ------------------------------------------------------------------------------------------------------------------
# rival routes that need more than a transfer basis

def at_risk_route(L, camps):
    """The houses still without PV after the campaign: each campaign's systems less the shortfall of its other first
    applications against the pre-year hazard (scaled by the comparison's change) on the houses left after the
    campaign's participants leave the pool."""
    out = []
    for c in camps:
        op = pd.Timestamp(c["opening"])
        end = C.window_end(op)
        p0 = op - pd.DateOffset(months=12)
        H = C.house_first(L)
        h = H[H.locality == c["suburb"].upper()]
        first, prog = h["first"], h["prog"].str.startswith("SS-")
        span = (end - op).days / 365.25
        comp = C.comparison(L, op, end, "nyt", c["suburb"])
        pre_h = C.n_between(first, p0, op) / C.without_pv(first, p0)
        cq = sum(C.n_between(C.suburb_first(L, s), p0, op) for s in comp) / sum(C.without_pv(C.suburb_first(L, s), p0) for s in comp)
        cp = sum(C.n_between(C.suburb_first(L, s), op, end) for s in comp) / sum(C.without_pv(C.suburb_first(L, s), op) for s in comp) / span
        organic = int(((first >= op) & (first < end) & ~prog).sum())
        left = C.without_pv(first, op) - c["systems"]
        expected = pre_h * cp / cq * left * span
        out.append(dict(c, net=c["systems"] - (expected - organic)))
    return out


def twfe_route(L):
    """A two-way fixed-effects regression on every suburb's yearly first applications per owner-occupied house without
    PV (years from September), with a single indicator for the years from a campaign's opening; the coefficient over
    three years on each 2026 suburb's houses without PV."""
    ops = C.openings(L)
    rows = []
    for s in list(ops) + list(C.NEVER):
        first = C.suburb_first(L, s)
        for y in range(2015, 2026):
            a, b = pd.Timestamp(y, 9, 1), pd.Timestamp(y + 1, 9, 1)
            risk = C.without_pv(first, a)
            op = ops.get(s)
            rows.append(dict(s=s, y=str(y), rate=C.n_between(first, a, b) / risk if risk else 0.0,
                             post=float(op is not None and op <= a + pd.Timedelta(days=10) and op <= pd.Timestamp(K.BASE))))
    pan = pd.DataFrame(rows)
    X = pd.get_dummies(pan[["s", "y"]], drop_first=True).astype(float)
    X["post"], X["const"] = pan.post, 1.0
    beta = np.linalg.lstsq(X.values, pan.rate.values, rcond=None)[0]
    return float(beta[list(X.columns).index("post")])


def ols_lift(camps, pools, y, cols):
    A = np.array([[1.0] + [float(c[k]) for k in cols] for c in camps])
    beta = np.linalg.lstsq(A, np.array([c[y] for c in camps], float), rcond=None)[0]
    return {n: float(beta @ np.array([1.0] + [float(p[k]) for k in cols])) for n, p in pools.items()}


def with_pen(rows):
    return [dict(r, pen=1.0 - r["pool"] / r["oo"]) for r in rows]


# ------------------------------------------------------------------------------------------------------------------
# the clean truth: the world before any file is written (dataset-generation 7.1)

def truth(W):
    """Connected capacity, trend, the FY2025-26 average, the seven campaigns (systems installed, houses without PV at
    the opening, households brought in from outside the market) and the 2026 suburbs' houses without PV, from the
    world's own records (no file, no join)."""
    nmi_feeder = {n["nmi"]: n["feeder"] for n in W.nmis}
    apps = list(W.apps)

    def cap(date):
        latest = {}
        for a in apps:
            if a["status"] in ("Connected", "Removed") and a["connected"] and a["connected"] <= date and \
                    not (a.get("removed") and a["removed"] <= date):
                k = (a["connected"], 0 if a["type"] == "Alteration" else 1)
                if a["nmi"] not in latest or k > latest[a["nmi"]][0]:
                    latest[a["nmi"]] = (k, a["kw"])
        out = collections.defaultdict(float)
        for n, (_, kw) in latest.items():
            if n in nmi_feeder:
                out[nmi_feeder[n]] += kw
        return out

    c0 = cap(K.BASE)
    yearly = collections.defaultdict(list)
    for s, e in K.TREND_FYS:
        c1, cb = cap(e), cap(s - dt.timedelta(days=1))
        prog = collections.defaultdict(float)
        for a in apps:
            if a["program"] and a["type"] == "New" and a["status"] in ("Connected", "Removed") and s <= a["connected"] <= e \
                    and a["nmi"] in nmi_feeder:
                prog[nmi_feeder[a["nmi"]]] += a["kw"]
        for f in K.FEEDERS:
            yearly[f].append(c1[f] - cb[f] - prog[f])
    trend = {f: float(np.mean(yearly[f])) for f in K.FEEDERS}
    base = {f: c0[f] + K.YEARS_TO_DESIGN * trend[f] for f in K.FEEDERS}
    s, e = K.LAST_FY
    by = collections.defaultdict(list)
    for a in apps:
        if a["type"] == "New" and a["status"] in ("Connected", "Removed") and s <= a["connected"] <= e and a["nmi"] in nmi_feeder:
            by[nmi_feeder[a["nmi"]]].append(a["kw"])
    avgkw = {f: float(np.mean(v)) for f, v in by.items()}
    first = {}
    for a in apps:
        if a["status"] != "Withdrawn" and (a["nmi"] not in first or a["received"] < first[a["nmi"]]):
            first[a["nmi"]] = a["received"]

    def suburb(name, opening):
        oo = [p for p in W.props if p["kind"] == "house_oo" and p["locality"] == name]
        nm = [p["nmis"][0]["nmi"] for p in oo]
        pool = sum(1 for n in nm if not (n in first and first[n] < opening))
        letter = sum(1 for n in W.nmis if n["locality"] == name and n["tariff"] == "Residential")
        return dict(oo=len(oo), pool=pool, letterboxed=letter, feeder=[l["feeder"] for l in K.LOCALITIES if l["name"] == name][0])

    camps = []
    for c in K.CAMPAIGNS:
        st = suburb(c["suburb"], c["open"])
        ref = "SS-%s-%s" % (c["key"][-2:], c["key"][:3])
        st.update(key=c["key"], systems=sum(1 for a in apps if a["program"] == ref and a["type"] == "New"
                                            and a["status"] in ("Connected", "Removed")),
                  net=float(W.camp_new[c["key"]]))
        camps.append(st)
    pools = {c["suburb"]: dict(suburb(c["suburb"], c["open"]), key=c["key"]) for c in K.CAMPAIGNS_2026}
    return dict(cap=c0, trend=trend, base=base, avgkw=avgkw, camps=camps, pools=pools)


def truth_route(T, L, measure, basis):
    rate = sum(c[measure] for c in T["camps"]) / sum(c[basis] for c in T["camps"])
    lift = {n: rate * p[basis] for n, p in T["pools"].items()}
    return with_base(L, T["base"], lift, T["avgkw"], T["pools"])


def fresh(L, **kw):
    """A copy of the loaded pack with the cached house tables dropped and some frames replaced."""
    out = {k: v for k, v in L.items() if not k.startswith("_houses_")}
    out.update(kw)
    return out


# ------------------------------------------------------------------------------------------------------------------

def run(target, golden, W=None, meta_path=None):
    N[0] = 0
    SF_SEEN.clear()
    L = C.load(target)
    reg = L["reg"]
    base = C.base_forecast(reg)
    camps = C.campaigns(L)
    pools = C.pools_2026(L)
    fc = C.forecast(L, base=base, camps=camps, pools=pools)
    ex = fc["excess"]
    files = sorted(os.listdir(target))
    npp = pdf_text(os.path.join(target, K.NPP))
    letter = pdf_text(os.path.join(target, K.LETTER))
    hca_txt = pdf_text(os.path.join(target, K.HCA))
    task_dir = os.path.dirname(os.path.abspath(target))
    prompt_txt = open(os.path.join(task_dir, "prompt.txt")).read()
    figs_path = os.path.join(task_dir, "generator", "figures.json")
    figs = json.load(open(figs_path)) if os.path.exists(figs_path) else None

    def fcast(camps_=None, **kw):
        return C.forecast(L, base=base, camps=camps_ if camps_ is not None else camps, pools=pools, **kw)

    def lifted(lift):
        return C.forecast(L, base=base, camps=camps, pools=pools, lift=lift)

    # ---------------------------------------------------------------- the call
    ok(graded(fc) == golden_tuple(), "golden %s" % (graded(fc),))
    ok(abs(r1(fc["gap"]) - (r1(ex["YRB27"]) - r1(ex["YRB24"]))) < 1e-9, "SF22 gap equals the rounded difference")
    ok(fc["gap"] >= 0.15, "golden margin %.3f" % fc["gap"])
    ok(ex["YRB24"] - ex[fc["order"][2]] >= 0.15, "the runner-up within 0.15 MW of third (%s)" % fc["order"][2])
    g0 = raw(fc)
    for k in FIGS:
        ok(bin_room(g0[k], 0.1) >= 0.30, "%s %.4f too near a rounding edge" % (k, g0[k]))
    if figs:
        cf = figs["call"]
        ok((cf["feeder"], cf["runner_up"], float(cf["excess_mw"]), float(cf["runner_up_excess_mw"]), float(cf["gap_mw"]),
            float(cf["bindarra_waters_mw"]), float(cf["wattleview_mw"])) == golden_tuple(), "figures.json call block")

    # ---------------------------------------------------------------- the filed WHAT, and no filed HOW for the programs
    ok("the impact on the feeder's connected capacity at the design point of each program" in npp, "NPP-114 4.1 estimand")
    low = npp.lower()
    for w in ("rate taken from past programs", "combined total", "letterbox", "first application", "uptake", "comparison",
              "not yet", "free rid", "additional", "displac", "pull", "counterfactual", "would have", "before and after",
              "per house", "per home", "systems installed", "the number of systems"):
        ok(w not in low, "NPP-114 names a method or basis: %r" % w)
    plow = prompt_txt.lower()
    for w in ("npp", "procedure", "net", "comparison", "would have", "anyway", "additional", "not yet", "before and after",
              "free rid", "pull", "control", "counterfactual"):
        ok(not re.search(r"\b%s" % re.escape(w), plow), "the prompt names a method: %r" % w)
    ok("how much each bulk-buy will raise its feeder's PV by that date" in prompt_txt, "the prompt's bulk-buy estimand")

    # ---------------------------------------------------------------- the decisive estimate: every defensible
    # construction of what a past campaign added lands every graded figure
    variants = {"not yet treated, 12 months before (golden)": camps}
    for mode in ("nyt", "never", "both"):
        for pre in (12, 24, 36):
            for form in ("ratio", "hazard"):
                key = "%s %s %d" % (form, mode, pre)
                if (mode, pre, form) != ("nyt", 12, "ratio"):
                    variants[key] = C.campaigns(L, pre=pre, mode=mode, form=form)
    for mode in ("nyt", "both"):
        for pre in (12, 24):
            variants["month by month %s %d" % (mode, pre)] = C.campaigns(L, pre=pre, mode=mode, form="cs")
    variants["24-month window"] = C.campaigns(L, months=24)
    for mode in ("nyt", "never", "both"):
        variants["by connection date %s" % mode] = C.campaigns(L, mode=mode, by="connection_date")
    variants["by connection date hazard"] = C.campaigns(L, form="hazard", by="connection_date")
    band = {}
    for name, cs in variants.items():
        v = fcast(cs)
        band[name] = v
        ok(graded(v) == golden_tuple(), "SF16/SF17/SF18/SF11/SF13 construction %s files %s" % (name, graded(v)))
    for est in ("mean", "median", "origin", "ols"):
        v = fcast(estimator=est)
        band["estimator " + est] = v
        ok(graded(v) == golden_tuple(), "SF15 estimator %s files %s" % (est, graded(v)))
    for drop in (("PBG24",), ("PBG24", "BNR23")):
        cs = [c for c in camps if c["key"] not in drop]
        v = C.forecast(L, base=base, camps=cs, pools=pools)
        band["without " + "+".join(drop)] = v
        ok(graded(v) == golden_tuple(), "SF25 without %s files %s" % (drop, graded(v)))
    for k in FIGS:
        vals = [raw(v)[k] for v in band.values()]
        lo, hi = min(vals), max(vals)
        ok(lo - (GOLD[k] - 0.05) >= 0.010 and (GOLD[k] + 0.05) - hi >= 0.010,
           "the constructions' corridor for %s %.3f to %.3f sits within 0.010 of a bin edge" % (k, lo, hi))
        SF_SEEN["corridor %s" % k] = (round(lo, 3), round(hi, 3))
    SF_SEEN["constructions"] = len(band)
    # one standard error of the pooled rate either way moves no graded figure and no feeder
    rr = C.rate_of(camps)
    se = C.ratio_se(camps)
    for sgn in (1, -1):
        v = lifted({n: (rr["slope"] + sgn * se) * p["pool"] for n, p in pools.items()})
        ok(graded(v) == golden_tuple(), "one standard error %+d moves the call: %s" % (sgn, graded(v)))
        SF_SEEN["one sigma %+d" % sgn] = tuple(round(raw(v)[k], 3) for k in FIGS)
    ok(0.005 <= se / rr["slope"] <= 0.03, "the pooled rate's standard error %.4f of %.5f" % (se, rr["slope"]))
    ok(0.02 <= rr["cv"] <= 0.10, "what the campaigns added per house without PV spreads %.3f" % rr["cv"])
    ok(abs(rr["slope"] * 1000 - round(rr["slope"] * 1000)) > 0.05, "a round pooled rate")
    SF_SEEN["pooled added per house"] = (round(rr["slope"], 5), round(se, 5), round(rr["cv"], 3))

    # ---------------------------------------------------------------- the past campaigns: what the records show
    gross = C.rate_of(camps, "systems")
    ops = C.openings(L)
    for c in camps:
        ok(c["systems_log"] == c["systems"], "%s log %d register %d" % (c["key"], c["systems_log"], c["systems"]))
        ok(c["letterboxed"] == c["res_nmis"], "%s letterboxed %d residential NMIs %d" % (c["key"], c["letterboxed"], c["res_nmis"]))
        ok(c["systems"] / c["net"] >= 1.7, "%s installed only %.2f times what it added" % (c["key"], c["systems"] / c["net"]))
        # the dip: the campaign suburb's other first applications fall well short of the comparison-implied level for
        # the two years after the opening, and come back after
        op = pd.Timestamp(c["opening"])
        H = C.house_first(L)
        h = H[H.locality == c["suburb"].upper()]
        prog = h["prog"].str.startswith("SS-")
        p0 = op - pd.DateOffset(months=12)
        comp = C.comparison(L, op, C.window_end(op), "nyt", c["suburb"])
        pre_c = sum(C.n_between(C.suburb_first(L, s), p0, op) for s in comp)

        def expected(a, b):
            return c["pre"] * sum(C.n_between(C.suburb_first(L, s), a, b) for s in comp) / pre_c
        mid = op + pd.DateOffset(months=24)
        dip = int(((h["first"] >= op) & (h["first"] < mid) & ~prog).sum()) / expected(op, mid)
        ok(dip <= 0.60, "%s other first applications ran at %.2f of the expected level in its first two years" % (c["key"], dip))
        end = C.window_end(op)
        if (end - mid).days >= 180:
            back = int(((h["first"] >= mid) & (h["first"] < end) & ~prog).sum()) / expected(mid, end)
            ok(0.85 <= back <= 1.15, "%s after two years at %.2f of the expected level" % (c["key"], back))
            SF_SEEN["dip %s" % c["key"]] = (round(dip, 2), round(back, 2))
        else:
            SF_SEEN["dip %s" % c["key"]] = (round(dip, 2), None)
    # the write-up's "more than half of the households the past bulk-buys signed up would have gone solar within three
    # years anyway": the seven together
    anyway = 1.0 - sum(c["net"] for c in camps) / sum(c["systems"] for c in camps)
    ok(0.53 <= anyway <= 0.60, "the seven installed %.3f more than they added" % anyway)
    SF_SEEN["share that would have gone solar anyway"] = round(anyway, 3)
    pp = C.campaigns(L, form="prepost")
    pp_rate = C.rate_of(pp)["slope"]
    ok(pp_rate / rr["slope"] >= 1.30, "each suburb's own before and after overstates by only %.2f" % (pp_rate / rr["slope"]))
    ratios = []
    for c in camps:
        op = pd.Timestamp(c["opening"])
        end = C.window_end(op)
        comp = C.comparison(L, op, end, "nyt", c["suburb"])
        span = (end - op).days / 365.25
        ratios.append(sum(C.n_between(C.suburb_first(L, s), op, end) for s in comp) /
                      sum(C.n_between(C.suburb_first(L, s), op - pd.DateOffset(months=12), op) for s in comp) / span)
    ok(max(ratios) / min(ratios) >= 1.15, "the comparison suburbs' change barely differs by campaign year")
    SF_SEEN["comparison yearly ratios"] = tuple(round(x, 3) for x in ratios)
    for name, p in pools.items():
        ok(p["letterboxed"] == p["res_nmis"], "%s letter %d NMIs %d" % (name, p["letterboxed"], p["res_nmis"]))
    hz = [c["n365"] / c["pool"] for c in camps]
    for name, p in pools.items():
        ok(min(hz) * 0.8 <= p["n365"] / p["pool"] <= max(hz) * 1.25, "%s's year-before uptake outside the past campaigns'" % name)
    pb, pw = pools["Bindarra Waters"], pools["Wattleview"]
    ok(min(c["pool"] for c in camps) < pw["pool"] < pb["pool"] < max(c["pool"] for c in camps),
       "the 2026 pools sit outside the past campaigns' range")
    SF_SEEN["installed per house"] = (round(gross["slope"], 5), round(gross["cv"], 3))

    # ---------------------------------------------------------------- the ladder and every rival route
    rungs = {"R0 installed systems per home letterboxed (default path)": fcast(measure="systems", basis="letterboxed"),
             "R1 installed systems per house without PV (the textbook transfer)": fcast(measure="systems", basis="pool")}
    for pre in (12, 24, 36):
        rungs["R2 each suburb's own before and after, %d months" % pre] = fcast(C.campaigns(L, pre=pre, form="prepost"))
    for name, v in rungs.items():
        ok(v["winner"] == "YRB24" and v["excess"]["YRB24"] - v["excess"]["YRB27"] >= 0.10,
           "%s files %s (YRB24 over YRB27 by %.3f)" % (name, v["winner"], v["excess"]["YRB24"] - v["excess"]["YRB27"]))
        ok(misses_all(v), "%s lands the call" % name)
        SF_SEEN["rung " + name] = (v["winner"], round(v["excess"][v["winner"]], 3), v["runner"],
                                   round(v["add_mw"]["Bindarra Waters"], 3), round(v["add_mw"]["Wattleview"], 3))
    r0 = rungs["R0 installed systems per home letterboxed (default path)"]
    ok(r0["runner"] == "YRB27" and r0["order"].index("YRB27") == 1, "R0's runner-up")
    if figs:
        ok(figs["default_path"]["feeder"] == "YRB24" and figs["reflex_path"]["feeder"] == "YRB24" and
           figs["prepost_path"]["feeder"] == "YRB24", "figures.json rungs")
    rivals = {"no campaigns": fcast(measure="none")}
    for b in ("count", "oo", "n365"):
        for est in (("pooled",) if b == "count" else ("pooled", "mean", "median", "ols")):
            rivals["installed per %s %s" % (b, est)] = fcast(measure="systems", basis=b, estimator=est)
    for b in ("letterboxed", "pool"):
        for est in ("mean", "median", "ols"):
            rivals["installed per %s %s" % (b, est)] = fcast(measure="systems", basis=b, estimator=est)
    for b in ("letterboxed", "oo", "n365", "count"):
        rivals["added per %s" % b] = fcast(measure="net", basis=b)
    for pre in (12, 24):
        rivals["every other suburb as the comparison, %d months" % pre] = fcast(C.campaigns(L, pre=pre, mode="all"))
    rivals["hazard of the houses left after the participants"] = fcast(at_risk_route(L, camps))
    beta = twfe_route(L)
    rivals["two-way fixed effects, three years"] = lifted({n: 3 * beta * p["pool"] for n, p in pools.items()})
    SF_SEEN["twfe coefficient"] = round(beta, 5)
    cpen, ppen = with_pen(camps), {n: dict(p, pen=1.0 - p["pool"] / p["oo"]) for n, p in pools.items()}
    for cols in (("pool",), ("letterboxed",), ("pool", "letterboxed"), ("oo", "pool"), ("pool", "pen"), ("letterboxed", "oo")):
        rivals["installed regressed on %s" % "+".join(cols)] = lifted(ols_lift(cpen, ppen, "systems", cols))
    for cols in (("letterboxed",), ("oo",), ("n365",)):
        rivals["added regressed on %s" % "+".join(cols)] = lifted(ols_lift(cpen, ppen, "net", cols))
    # a regression of what was added on the houses without PV, with an intercept, or with penetration beside it, is
    # the golden construction again
    for cols in (("pool",), ("pool", "pen")):
        v = lifted(ols_lift(cpen, ppen, "net", cols))
        ok(graded(v) == golden_tuple(), "SF15 added regressed on %s files %s" % ("+".join(cols), graded(v)))
        SF_SEEN["SF15 added regressed on %s" % "+".join(cols)] = tuple(round(raw(v)[k], 3) for k in FIGS)
    guesses = {}
    for share in (0.3, 1.0 / 3, 0.4, 0.5, 0.6, 0.7):
        guesses["an assumed %.0f%% of installed systems additional" % (100 * share)] = lifted(
            {n: share * gross["slope"] * p["pool"] for n, p in pools.items()})
    guesses["installed less a year of the suburb's own uptake"] = lifted(
        {n: gross["slope"] * p["pool"] - p["n365"] for n, p in pools.items()})
    for name, v in guesses.items():
        g = graded(v)
        ok(v["winner"] != GOLD["winner"] or (g[2] != GOLD["excess"] and g[5] != GOLD["bdw"] and g[6] != GOLD["wtv"]),
           "guess %s lands the excess or a bulk-buy figure: %s" % (name, g))
        SF_SEEN["guess " + name] = g
    for name, v in rivals.items():
        ok(misses_all(v), "rival %s lands the call: %s" % (name, graded(v)))
        ok(not (v["runner"] == GOLD["runner"] and r1(v["excess"][GOLD["runner"]]) == GOLD["runner_excess"] and
                v["winner"] == GOLD["winner"]), "rival %s files the golden's runner-up and its figure" % name)
        SF_SEEN["rival " + name] = graded(v)
    # the five routes the rework names: each lands another feeder or misses every figure by a bin or more
    named = {"gross per campaign": rivals["installed per count pooled"], "per home letterboxed": r0,
             "per eligible house": rungs["R1 installed systems per house without PV (the textbook transfer)"],
             "pooled before and after": rungs["R2 each suburb's own before and after, 12 months"],
             "fitted on the obvious covariates": rivals["installed regressed on pool+letterboxed"]}
    for name, v in named.items():
        g = raw(v)
        ok(v["winner"] != GOLD["winner"] or all(abs(g[k] - g0[k]) >= 0.1 for k in FIGS),
           "%s within a bin of the golden" % name)
    # the assumed 45 per cent share sits near the pooled rate's share (recorded, a guess the records do not support)
    v45 = lifted({n: 0.45 * gross["slope"] * p["pool"] for n, p in pools.items()})
    SF_SEEN["an assumed 45% additional (recorded)"] = graded(v45)
    SF_SEEN["pooled share additional"] = round(rr["slope"] / gross["slope"], 3)

    # ---------------------------------------------------------------- the three rules, executed (dataset-generation 7.1)
    heads = " ".join(open(os.path.join(target, f), encoding="utf-8").readline().lower() for f in files if f.endswith(".csv"))
    ok(not re.search(r"net_|additional|free_rid|uplift|counterfactual|expected|displac|first_app|participation", heads),
       "R1: a shipped column holds the deciding quantity")
    ok({"received_date", "status", "application_type", "nmi", "program_ref", "connection_date"} <= set(reg.columns) and
       {"land_use_code", "rating_category", "locality"} <= set(L["rates"].columns) and len(L["xref"]) > 50000,
       "R1: the register's application rows, the rates roll and the cross-reference ship at their own grain")
    import openpyxl
    wbp = openpyxl.load_workbook(os.path.join(target, K.PROGRAM_LOG), read_only=True)
    plog_txt = " ".join(str(c) for ws in wbp.worksheets for row in ws.iter_rows(values_only=True) for c in row if c)
    ok(not re.search(r"anyway|additional|would have|per cent of|participation rate|expect|net of", (plog_txt + " " + letter).lower()),
       "R1: the council's files state a rate, an expectation or an effect")
    if W is not None:
        T = truth(W)
        # R2: the mess on the call's path (application grain with alterations, withdrawn and unconnected rows, removals,
        # migrated rows with inverters in the legacy table, feeder at application after load transfers, the rates roll
        # joined through the GIS cross-reference, the council's workbook layout), read by the filed rules and the
        # dictionary, gives back the world's own figures
        for f in K.FEEDERS:
            ok(abs(base["cap"].get(f, 0.0) - T["cap"][f]) < 1e-6, "R2: %s connected capacity %.3f, truth %.3f"
               % (f, base["cap"].get(f, 0.0), T["cap"][f]))
            ok(abs(base["trend"][f] - T["trend"][f]) < 1e-6, "R2: %s trend differs from the truth" % f)
        ok(all(abs(fc["avgkw"][f] - T["avgkw"][f]) < 1e-9 for f in K.FEEDERS), "R2: the FY2025-26 averages differ")
        for c, t in zip(camps, T["camps"]):
            ok((c["systems"], c["pool"], c["letterboxed"], c["oo"]) == (t["systems"], t["pool"], t["letterboxed"], t["oo"]),
               "R2: %s read from the pack differs from the truth" % c["key"])
            ok(t["pool"] == W.camp_pool[c["key"]] and t["systems"] == W.camp_systems[c["key"]],
               "R2: %s the simulation's own counts" % c["key"])
        for name, p in pools.items():
            t = T["pools"][name]
            ok((p["pool"], p["letterboxed"], p["oo"]) == (t["pool"], t["letterboxed"], t["oo"]),
               "R2: %s read from the pack differs from the truth" % name)
        # each campaign's estimate against what it truly added (households brought in from outside the market)
        for c, t in zip(camps, T["camps"]):
            ok(abs(c["net"] - t["net"]) <= max(6.0, 0.07 * t["net"]), "R2: %s estimate %.1f against the truth %.0f"
               % (c["key"], c["net"], t["net"]))
        # R3a: the easy route on the clean truth (installed systems per home letterboxed, every filed rule applied)
        # files the default path's wrong answer; the textbook transfer of installed systems files it too
        t0 = truth_route(T, L, "systems", "letterboxed")
        ok(graded(t0) == graded(r0) and t0["winner"] == "YRB24", "R3a: the per-letterbox route on the clean truth files %s"
           % (graded(t0),))
        t1 = truth_route(T, L, "systems", "pool")
        ok(t1["winner"] == "YRB24" and graded(t1) == graded(rungs["R1 installed systems per house without PV (the textbook transfer)"]),
           "R3a: installed systems per house without PV on the clean truth files %s" % (graded(t1),))
        # R3b: the same route with only the decisive step changed (what each campaign added, not what it installed)
        # files the golden
        tg = truth_route(T, L, "net", "pool")
        ok(graded(tg) == golden_tuple(), "R3b: what the campaigns added on the clean truth files %s" % (graded(tg),))
        ok(truth_route(T, L, "net", "letterboxed")["winner"] != "YRB27" or misses_all(truth_route(T, L, "net", "letterboxed")),
           "R3b: the added systems per home letterboxed on the clean truth land the call")
        SF_SEEN["R3 truth"] = (graded(t0), graded(t1), graded(tg))
    # R3c: nothing planted
    ok(not any(w in f.lower() for f in files for w in ("memo", "draft", "briefing", "estimate", "evaluation")),
       "R3c: a memo, draft, briefing, estimate or evaluation ships")
    docs_txt = " ".join(pdf_text(os.path.join(target, f)) for f in files if f.endswith(".pdf"))
    ok(not re.search(r"\b\d+ systems\b|\bsystems per\b|\bper cent of (?:the )?(?:homes|houses|households)\b|anyway|"
                     r"additional|would have", docs_txt), "R3c: a filed document states a campaign forecast, a rate or an effect")
    ok("excess" not in hca_txt.lower() and "headroom" not in hca_txt.lower() and "rank" not in hca_txt.lower(),
       "R3c: the assessment ranks the feeders or states an excess")

    # ---------------------------------------------------------------- the semantic fork rows
    def variant(base_kw=None, avgkw=None, lift=None, camps_=None, pools_=None, hc_case="hc29"):
        cc = camps_ or camps
        pp = pools_ or pools
        lf = lift if lift is not None else C.lifts(cc, pp)
        return with_base(L, base_kw if base_kw is not None else base["base"], lf,
                         avgkw if avgkw is not None else fc["avgkw"], pp, hc_case)

    # SF1 residential NMIs only
    res = set(L["net"][L["net"].tariff_class == "Residential"].nmi)
    reg_res = reg[reg.nmi.isin(res)]
    v = variant(base_kw=base_with(reg_res, C.connected_kw, C.trend(reg_res)[0]))
    ok(graded(v) != golden_tuple(), "SF1 residential only keeps every figure")
    SF_SEEN["SF1 residential only"] = graded(v)
    ok("at a residential or business NMI" in npp, "SF1 NPP-114 2.1 no longer names business NMIs")
    ok(reg.kw.max() <= 30.0, "SF1 a register system above 30 kVA (%.1f)" % reg.kw.max())
    # SF2 alterations summed; panel capacity
    tr2 = {}
    for f in K.FEEDERS:
        vals = []
        for s, e in K.TREND_FYS:
            c1, c0 = summed_capacity(reg, e), summed_capacity(reg, s - dt.timedelta(days=1))
            vals.append(c1.get(f, 0) - c0.get(f, 0) - C.program_kw(reg, s, e).get(f, 0))
        tr2[f] = float(np.mean(vals))
    v = variant(base_kw=base_with(reg, summed_capacity, tr2))
    ok(graded(v) != golden_tuple(), "SF2 summed alterations keeps every figure")
    ok("An alteration replaces the configuration it alters." in npp, "SF2 NPP-114 2.2 alteration sentence missing")
    SF_SEEN["SF2 summed"] = graded(v)
    reg_p = reg.assign(kw=reg.panel_kwp.replace("", np.nan).astype(float))
    v = variant(base_kw=base_with(reg_p, C.connected_kw, C.trend(reg_p)[0]), avgkw=C.avg_new_kw(reg_p))
    ok(graded(v) != golden_tuple(), "SF2 panel capacity keeps every figure")
    ok("inverter capacity (kW, AC rating)" in npp, "SF2 NPP-114 2.2 inverter capacity sentence missing")
    SF_SEEN["SF2 panel"] = graded(v)
    # SF3 migrated rows dropped
    reg_m = reg[reg.legacy_application_no == ""]
    v = variant(base_kw=base_with(reg_m, C.connected_kw, C.trend(reg_m)[0]))
    SF_SEEN["SF3 migrated dropped"] = graded(v)
    ok(graded(v) != golden_tuple(), "SF3 dropping the migrated rows keeps every figure")
    kw_def = [r_[1] for r_ in openpyxl.load_workbook(os.path.join(target, K.DICTIONARY), read_only=True)
              ["Connections register"].iter_rows(values_only=True) if r_ and r_[0] == "inverter_kw_ac"]
    ok(len(kw_def) == 1 and "Blank on applications migrated from the legacy system." in kw_def[0],
       "SF3 the dictionary no longer says where the migrated inverters are")
    reg_pk = reg.assign(kw=np.where(reg.legacy_application_no != "", reg.panel_kwp.replace("", np.nan).astype(float), reg.kw))
    v = variant(base_kw=base_with(reg_pk, C.connected_kw, C.trend(reg_pk)[0]))
    SF_SEEN["SF3 migrated at panel capacity"] = graded(v)
    # SF23 the legacy join is one to one
    leg = L["leg"]
    regl = reg[reg.legacy_application_no != ""]
    ok(leg.legacy_application_no.is_unique and regl.legacy_application_no.is_unique and
       set(regl.legacy_application_no) == set(leg.legacy_application_no), "SF23 the legacy numbers do not pair one to one")
    # SF4 approved but not connected counted
    reg4 = reg.copy()
    m = reg4.status.isin(["Approved"])
    reg4.loc[m, "status"] = "Connected"
    reg4.loc[m, "connection_date"] = reg4.loc[m, "approved_date"]
    b4 = base_with(reg4, C.connected_kw, C.trend(reg)[0])
    v = variant(base_kw=b4)
    ok(v["winner"] == "YRB27", "SF4 files %s" % v["winner"])
    ok(max(abs(b4[f] / base["base"][f] - 1) for f in K.FEEDERS) < 0.02, "SF4 moves a base by 2 per cent or more")
    SF_SEEN["SF4 approved counted"] = graded(v)
    # SF5 the feeder recorded on the application
    reg5 = reg.assign(feeder=reg.feeder_at_application)
    v = variant(base_kw=base_with(reg5, C.connected_kw, C.trend(reg5)[0]))
    ok(v["winner"] == "YRB27", "SF5 feeder at application files %s" % v["winner"])
    SF_SEEN["SF5 feeder at application"] = graded(v)
    # SF6 base date 30 June 2026
    v = variant(base_kw=base_with(reg, C.connected_kw, base_date=dt.date(2026, 6, 30)))
    ok(v["winner"] == "YRB27", "SF6 files %s" % v["winner"])
    SF_SEEN["SF6 30 June"] = graded(v)
    # SF7 trend variants
    tr_new = {}
    for f in K.FEEDERS:
        vals = []
        for s, e in K.TREND_FYS:
            new = reg[(reg.application_type == "New") & reg.status.isin(C.CONNECTED) & (reg.connection_date >= pd.Timestamp(s))
                      & (reg.connection_date <= pd.Timestamp(e)) & (reg.program_ref == "") & (reg.feeder == f)]
            vals.append(new.kw.sum())
        tr_new[f] = float(np.mean(vals))
    for name, trv in (("new connections only", tr_new), ("program systems left in", C.trend(reg, exclude_program=False)[0])):
        v = variant(base_kw=base_with(reg, C.connected_kw, trv))
        ok(v["winner"] == "YRB27", "SF7 %s files %s" % (name, v["winner"]))
        SF_SEEN["SF7 " + name] = graded(v)
    ok("less the capacity of the program systems first connected in that year" in npp, "SF7 NPP-114 3.2 sentence")
    # SF8 3.25 or 2.92 years
    for yrs in (3.25, 2.92):
        v = variant(base_kw=base_with(reg, C.connected_kw, years=yrs))
        ok(v["winner"] == "YRB27", "SF8 %.2f years files %s" % (yrs, v["winner"]))
        SF_SEEN["SF8 %.2f years" % yrs] = graded(v)
    # SF9 a campaign's installed systems: the program-referenced connected new systems, equal to the council's
    # systems installed (asserted above); the golden counts every first application, program or not, so the contracts
    # column enters no golden figure
    contracts = {r_["Suburb"]: int(r_["Contracts signed"]) for r_ in L["log"]}
    ok(all(contracts[c["suburb"]] > c["systems"] for c in camps), "SF9 contracts signed at or below systems installed")
    v = variant(lift={n: C.rate_of([dict(c, systems=contracts[c["suburb"]]) for c in camps], "systems")["slope"] * p["pool"]
                      for n, p in pools.items()})
    ok(v["winner"] == "YRB24", "SF9 contracts per house without PV files %s" % v["winner"])
    SF_SEEN["SF9 contracts installed per house"] = graded(v)
    # SF10 an added system's size: the feeder's FY2025-26 average new inverter (NPP-114 4.3), against the past program
    # systems' own average
    prog_kw = sum(c["kw_program"] for c in camps) / sum(c["systems"] for c in camps)
    v = variant(avgkw={f: prog_kw for f in K.FEEDERS})
    ok(graded(v) != golden_tuple() and v["winner"] == "YRB27", "SF10 the program systems' size files %s" % (graded(v),))
    SF_SEEN["SF10 past program size %.2f kW" % prog_kw] = graded(v)
    ok("taken at the average inverter capacity of the micro EG systems first connected on the feeder in the last "
       "complete financial year" in npp, "SF10 NPP-114 4.3 sentence missing")
    for name, p in pools.items():
        ok(abs(fc["add_mw"][name] - fc["lift"][name] * fc["avgkw"][p["feeder"]] / 1000.0) < 1e-12, "SF10 %s" % name)
    # SF12 whose first applications: every house (owner-occupied and rented, land use 02), or every residential address
    rates = L["rates"]
    Hall = rates[rates.land_use_code == "02"]
    ok(len(Hall) > len(rates[(rates.land_use_code == "02") & (rates.rating_category == "R1")]), "SF12 no rented houses")
    for scope, keep in (("every house", lambda r_: r_.land_use_code == "02"),
                        ("every residential address", lambda r_: r_.land_use_code.isin(["02", "08", "09"]))):
        rr_ = rates[keep(rates)]
        L2 = fresh(L, rates=rr_.assign(rating_category="R1", land_use_code="02"))
        cs2 = C.campaigns(L2)
        ps2 = C.pools_2026(L2)
        v = C.forecast(L2, base=base, camps=cs2, pools=ps2)
        SF_SEEN["SF12 " + scope] = graded(v)
        ok(misses_all(v) or graded(v) == golden_tuple(), "SF12 %s lands part of the call: %s" % (scope, graded(v)))
    terms = " ".join(str(c) for row in openpyxl.load_workbook(os.path.join(target, K.PROGRAM_LOG), read_only=True)
                     ["Program terms"].iter_rows(values_only=True) for c in row if c)
    ok("owner-occupiers of houses" in terms and "Tenants" in terms, "SF12 the program terms")
    # SF14 a suburb by the network model's locality: the same NMIs
    H = C.house_first(L)
    for name in list(ops):
        nm = set(L["xref"][L["xref"].property_no.isin(H[H.locality == name.upper()].index)].nmi)
        net_loc = set(L["net"][L["net"].locality.str.upper() == name.upper()].nmi)
        ok(nm <= net_loc, "SF14 %s rates-roll houses outside the network model's locality" % name)
    # SF19 how much a bulk-buy raises its feeder's PV by the design point: with the campaign against without it; the
    # systems it installs (R0, R1) are refused by the prompt's "raise its feeder's PV by that date" and NPP-114 4.1
    ok("raise its feeder's PV by that date" in prompt_txt and "impact on the feeder's connected capacity at the design "
       "point" in npp, "SF19 the estimand's sentences")
    # SF20 the transfer basis: per house without PV at the opening; per home letterboxed, per owner-occupied house and
    # per first application in the year before each scatter more across the seven and file another call or figures
    for b in ("letterboxed", "oo", "n365"):
        rb = C.rate_of(camps, "net", b)
        ok(rb["cv"] >= 1.5 * rr["cv"], "SF20 added per %s spreads %.3f against %.3f" % (b, rb["cv"], rr["cv"]))
        SF_SEEN["SF20 added per %s cv" % b] = round(rb["cv"], 3)
    # SF26 how much PV a bulk-buy adds: its systems at the feeder's FY2025-26 average new inverter, in MW (the prompt's
    # "megawatts to one decimal"; NPP-114 4.3 sizes every system not connected at the base date), never a count
    ok("Give megawatts to one decimal" in prompt_txt and "Systems a forecast counts that are not connected at the base "
       "date are taken at the average inverter capacity" in npp, "SF26 the unit and the sizing sentences")
    ok(all(abs(fc["add_mw"][n] * 1000.0 - fc["lift"][n] * fc["avgkw"][p["feeder"]]) < 1e-9 for n, p in pools.items()),
       "SF26 a bulk-buy's MW is not its systems at the feeder's average")
    # SF21 the 2026 case
    v = fcast(hc_case="hc26")
    ok(graded(v) != golden_tuple(), "SF21 the 2026 case keeps every figure")
    ok("the figure in the current hosting capacity assessment for the design year's spring minimum-demand case" in npp
       and "Design point: 31 August of the design year" in npp, "SF21 NPP-114 5.1 or 2.4 sentence missing")
    SF_SEEN["SF21 the 2026 case"] = graded(v)
    # SF24 a withdrawn application read as a house's first application
    regW = reg.assign(status=reg.status.replace({"Withdrawn": "Connected"}))
    L3 = fresh(L, reg=regW)
    cs3, ps3 = C.campaigns(L3), C.pools_2026(L3)
    v = C.forecast(L3, base=base, camps=cs3, pools=ps3)
    SF_SEEN["SF24 withdrawn counted"] = graded(v)
    ok(v["winner"] == "YRB27", "SF24 withdrawn as a first application files %s" % v["winner"])
    wbd = openpyxl.load_workbook(os.path.join(target, K.DICTIONARY), read_only=True)
    sdef = [r_[1] for r_ in wbd["Connections register"].iter_rows(values_only=True) if r_ and r_[0] == "status"]
    ok(len(sdef) == 1 and "nothing was connected under the application." in sdef[0], "SF24 the dictionary's Withdrawn")
    # SF29 a house with two NMIs takes the earlier: no campaign suburb has one
    xr = L["xref"].merge(L["rates"][["property_no", "locality"]], on="property_no")
    multi = xr.groupby("property_no").nmi.count()
    sub = {s.upper() for s in ops}
    ok(not xr[xr.property_no.isin(multi[multi > 1].index)].locality.isin(sub).any(), "SF29 dual occupancy in a campaign suburb")

    # ---------------------------------------------------------------- closure axes
    pend = reg[(reg.program_ref != "") & reg.status.isin(["Approved", "Under assessment"])]
    ok(len(pend) == 0, "closure 12: program systems pending at the base date")
    ok(not reg.program_ref.str.startswith("SS-26").any(), "closure 22: a 2026 program reference already lodged")
    for c in camps:
        ref = "SS-%s-%s" % (c["key"][-2:], c["key"][:3])
        conn = reg[(reg.program_ref == ref) & reg.status.isin(C.CONNECTED)].connection_date
        if pd.Timestamp(c["opening"]) >= pd.Timestamp(K.TREND_FYS[0][0]):
            ok(len({K.fy_of(x.date()) for x in conn}) == 1, "closure 3: %s program systems span financial years" % c["key"])
        ok((conn - pd.Timestamp(c["opening"])).dt.days.max() <= 270, "closure 3: %s connection lag" % c["key"])
    g0t = golden_tuple()
    ok((reg.connection_date == pd.Timestamp(K.BASE)).sum() > 0, "closure 9: no connection on the base date to test")
    v = variant(base_kw=base_with(reg, C.connected_kw, base_date=K.BASE - dt.timedelta(days=1)))
    ok(graded(v) == g0t, "closure 9: base date read exclusively")
    late = [(s_ + dt.timedelta(days=1), e_ + dt.timedelta(days=1)) for s_, e_ in K.TREND_FYS]
    v = variant(base_kw=base_with(reg, C.connected_kw, C.trend(reg, fys=late)[0]))
    ok(graded(v) == g0t, "closure 9: financial years a day late")
    # the campaign windows read a day late at both ends (records on the edges move to the next window)
    on_edge = 0
    for c in camps:
        op = pd.Timestamp(c["opening"])
        f_ = C.suburb_first(L, c["suburb"])
        on_edge += int(f_.eq(op).sum() + f_.eq(op - pd.DateOffset(months=12)).sum())
    camps_late = []
    for c in camps:
        e_ = C.campaign_effect(L, c["suburb"], pd.Timestamp(c["opening"]) + pd.Timedelta(days=1))
        camps_late.append(dict(c, net=e_["net"], pool=e_["pool"]))
    v = fcast(camps_late)
    ok(graded(v) == g0t, "closure 9: every campaign window a day late files %s" % (graded(v),))
    SF_SEEN["closure 9 records on a window edge"] = on_edge
    shuffled = fresh(L, reg=reg.sample(frac=1.0, random_state=7).reset_index(drop=True))
    v = C.forecast(shuffled, base=C.base_forecast(shuffled["reg"]), camps=C.campaigns(shuffled), pools=C.pools_2026(shuffled))
    ok(abs(v["excess"]["YRB27"] - ex["YRB27"]) < 1e-9, "closure 14 row order")
    ok(reg.nmi.str.len().eq(10).all() and L["net"].nmi.str.len().eq(10).all() and L["xref"].nmi.str.len().eq(10).all(),
       "closure 16 NMI formats on the main path")
    ok(ex[fc["order"][0]] - ex[fc["order"][1]] > 0.1, "closure 11")
    for c in K.CAMPAIGNS_2026:
        ok(c["open"] == K.BASE + dt.timedelta(days=1), "closure 4: a 2026 opening is not the day after the base date")
    ok("1 September 2026" in letter and "27 November 2026" in letter, "closure 22 the letter's sign-up dates")
    for c in camps:
        ok(pd.Timestamp(c["opening"]) == pd.Timestamp(K.CAMP[c["key"]]["open"]), "closure 4: %s opening" % c["key"])

    # ---------------------------------------------------------------- the asks
    asks = {"D sites": (C.ask_d, ("stamp", "connections", "no_map", "raw"), False),
            "A first lodged": (C.ask_a, ("active_first", "no_map", "raw"), False),
            "B storage": (C.ask_b, ("once", "rows", "stamp", "no_map"), True),
            "C inverter": (C.ask_c, ("once", "rows", "dedupe", "stamp", "no_map"), True)}
    primary = {"D sites": ("stamp",), "A first lodged": ("active_first",), "B storage": ("once", "rows"),
               "C inverter": ("once", "rows", "dedupe")}
    gold_asks = {}
    for name, (fn, modes, tens) in asks.items():
        g = fn(L)
        gold_asks[name] = g
        key = ten if tens else (lambda x: x)
        for m in modes:
            v = fn(L, m)
            hits = sum(key(v[f]) == key(g[f]) for f in K.FEEDERS)
            if m in primary[name]:
                ok(hits == 0, "ask %s primary stop %s lands %d feeders" % (name, m, hits))
            else:
                ok(hits <= 5, "ask %s hazard %s lands %d feeders" % (name, m, hits))
            SF_SEEN["ask %s %s misses" % (name, m)] = 8 - hits
        if tens:
            for f in K.FEEDERS:
                ok(bin_room(g[f], 10) >= 0.07, "ask %s %s %.2f too near a rounding edge" % (name, f, g[f]))
        ok(max(collections.Counter(g.values()).values()) <= 2, "ask %s: three feeders share a figure" % name)
    ok(all(C.ask_a(L, "stamp")[f] == gold_asks["A first lodged"][f] for f in K.FEEDERS),
       "ask A: the feeder stamp moves a site first lodged after the last transfer")
    # SF27 the feeder that supplies a DER site: the network model's, never the extract's stamp (the prompt's "it
    # supplies"; the network model's feeder is the feeder supplying the NMI at 31 August 2026, the extract's is only
    # "added by Mirrabee's extraction"); SF28 a site first lodged: the earliest lodgement at the site, decommissioned
    # connections included (the prompt counts sites, "how many of them"; the register dates each AC connection)
    ok("how many active DER Register sites it supplies" in prompt_txt and "how many of them were first lodged" in prompt_txt,
       "SF27 and SF28 the prompt's wording")
    wbd_ = openpyxl.load_workbook(os.path.join(target, K.DICTIONARY), read_only=True)
    ddef = {r_[0]: r_[1] for r_ in wbd_["DER Register extract"].iter_rows(values_only=True) if r_ and r_[0]}
    ndef = {r_[0]: r_[1] for r_ in wbd_["Network model"].iter_rows(values_only=True) if r_ and r_[0]}
    ok(ndef.get("feeder") == "Feeder supplying the NMI at 31 August 2026." and
       ddef.get("feeder") == "Feeder, added by Mirrabee's extraction.", "SF27 the dictionary's two feeders")
    ok(ddef.get("registered_date") == "Date the installer lodged the AC connection." and
       ddef.get("ac_connection_status") == "Active or Decommissioned.", "SF28 the dictionary's lodgement and status")
    st_ = C.ask_d(L, "stamp")
    ok(sum(st_[f] != gold_asks["D sites"][f] for f in K.FEEDERS) == 8, "SF27 the stamp lands a feeder")
    af_ = C.ask_a(L, "active_first")
    ok(sum(af_[f] != gold_asks["A first lodged"][f] for f in K.FEEDERS) == 8, "SF28 active-only lands a feeder")
    SF_SEEN["SF27 the extract's feeder"] = st_
    SF_SEEN["SF28 active connections only"] = af_
    if figs:
        for f in K.FEEDERS:
            d = figs["der"][f]
            ok((d["active_sites"], d["first_lodged_fy2025_26"], int(d["inverter_kw"].replace(",", "")),
                int(d["storage_kwh"].replace(",", ""))) == (gold_asks["D sites"][f], gold_asks["A first lodged"][f],
                                                             ten(gold_asks["C inverter"][f]), ten(gold_asks["B storage"][f])),
               "figures.json DER row %s" % f)
    # the asks' answers against the world
    if W is not None:
        tsites, tfirst = collections.Counter(), collections.Counter()
        tinv, tbat = collections.defaultdict(float), collections.defaultdict(float)
        by_site = collections.defaultdict(list)
        for s in W.der_systems:
            by_site[s["perm"]].append(s)
        nmi_by = {n["nmi"]: n for n in W.nmis}
        for perm, ss in by_site.items():
            act_s = [s for s in ss if s["status"] == "Active"]
            if not act_s or perm not in nmi_by:
                continue
            f = nmi_by[perm]["feeder"]
            tsites[f] += 1
            if K.LAST_FY[0] <= min(s["registered"] for s in ss) <= K.LAST_FY[1]:
                tfirst[f] += 1
            for s in act_s:
                tinv[f] += s["kw"]
                tbat[f] += s["storage"]
        ok(all(tsites[f] == gold_asks["D sites"][f] for f in K.FEEDERS), "ask D against the world")
        ok(all(tfirst[f] == gold_asks["A first lodged"][f] for f in K.FEEDERS), "ask A against the world")
        ok(all(abs(tinv[f] - gold_asks["C inverter"][f]) < 1e-6 for f in K.FEEDERS), "ask C against the world")
        ok(all(abs(tbat[f] - gold_asks["B storage"][f]) < 1e-6 for f in K.FEEDERS), "ask B against the world")
    der = L["der"]
    act = der[der.ac_connection_status == "Active"]
    keys = C.der_keys(L)
    keys = keys.assign(net=keys.key.map(L["feeder"]))
    act_k = keys[keys.ac_connection_status == "Active"]
    moved = act_k[act_k.net.notna() & (act_k.net != act_k.feeder)]
    ok(set(moved.net) | set(moved.feeder) == set(K.FEEDERS), "ask D stamp device misses a feeder: %s"
       % sorted(set(moved.net) | set(moved.feeder)))
    SF_SEEN["ask D stamp rows"] = (len(moved), moved.ac_connection_id.nunique())
    bat = act[act.device_type == "Battery"]
    sbr = act_k[(act_k.device_type == "Battery") & act_k.model.isin(C.SYSTEM_LODGED)]
    ok(set(sbr.net.dropna()) == set(K.FEEDERS), "ask B: a feeder without a Sungrow SBR stack")
    modular = bat[bat.model.isin(["Battery-Box Premium HVS", "Lynx Home U"])]
    ok(len(modular) > 0, "ask B: no module-by-module battery")
    inv = act[act.device_type == "Inverter"].assign(kw=lambda d: d.rated_power_kw.astype(float))
    g_inv = inv.groupby("ac_connection_id").agg(n=("kw", "size"), kw=("kw", "first"), model=("model", "first"),
                                                reg=("registered_date", "first"), nmi=("nmi", "first"))
    duals = g_inv[g_inv.n > 1].assign(plate=lambda d: d.model.map(C.nameplate))
    old_d = duals[duals.reg < K.iso(K.PORTAL_SWITCH)]
    new_d = duals[duals.reg >= K.iso(K.PORTAL_SWITCH)]
    ok(old_d.plate.notna().all() and new_d.plate.notna().all(), "ask C: a two-inverter connection without a nameplate")
    ok((abs(old_d.kw - old_d.plate * old_d.n) < 0.011).all() and (abs(new_d.kw - new_d.plate) < 0.011).all(),
       "ask C: a two-inverter row's rating against its nameplate")
    ok(len(old_d) >= 40 and len(new_d) >= 8, "ask C devices %d old, %d new two-inverter connections" % (len(old_d), len(new_d)))
    SF_SEEN["ask C duals"] = (len(old_d), len(new_d))
    # the referee: the register's own figures for each NMI agree with the golden handling, never with the stops
    remap = dict(zip(L["realloc"].old_nmi, L["realloc"].new_nmi))
    live = reg[reg.status.isin(C.CONNECTED) & (reg.connection_date <= pd.Timestamp(K.BASE)) &
               (reg.removal_date.isna() | (reg.removal_date > pd.Timestamp(K.BASE)))]
    last = live.sort_values(["nmi", "connection_date", "application_type"]).groupby("nmi").tail(1).set_index("nmi")
    bt = bat.assign(kwh=bat.storage_capacity_kwh.astype(float), key=keys.loc[bat.index, "key"])
    per = bt.groupby("ac_connection_id").agg(k=("kwh", "size"), v=("kwh", "first"), s=("kwh", "sum"),
                                             model=("model", "first"), key=("key", "first"))
    per = per.assign(app=per.key.map(last.battery_kwh).replace("", np.nan).astype(float)).dropna(subset=["app"])
    sb, ot = per[per.model.isin(C.SYSTEM_LODGED)], per[~per.model.isin(C.SYSTEM_LODGED)]
    ok(len(sb) >= 20 and (abs(sb.v - sb.app) < 0.011).mean() >= 0.9 and (abs(sb.s - sb.app) < 0.011).mean() <= 0.05,
       "ask B referee: SBR rows against the register's battery_kwh")
    ok((abs(ot.s - ot.app) < 0.011).mean() >= 0.9, "ask B referee: module rows against the register's battery_kwh")
    gi = g_inv.assign(key=g_inv.nmi.map(C.norm_nmi).map(lambda x: remap.get(x, x)))
    gi = gi.assign(app=gi.key.map(last.kw)).dropna(subset=["app"])
    od, nd = gi[(gi.n > 1) & (gi.reg < K.iso(K.PORTAL_SWITCH))], gi[(gi.n > 1) & (gi.reg >= K.iso(K.PORTAL_SWITCH))]
    ok((abs(od.kw - od.app) < 0.011).mean() >= 0.8 and (abs(nd.kw * nd.n - nd.app) < 0.011).mean() >= 0.8,
       "ask C referee: two-inverter rows against the register's inverter_kw_ac")
    # zero device rows in the main call's population: the call ignores the DER extract and the re-allocation log
    ok(graded(C.forecast(fresh(L, der=L["der"].iloc[0:0], realloc=L["realloc"].iloc[0:0]), base=base, camps=camps, pools=pools)) == golden_tuple(), "R2: the call reads an ask file")
    ok(not reg.nmi.isin(set(L["realloc"].old_nmi)).any(), "a replaced NMI in the connection register")

    # ---------------------------------------------------------------- distractors: deleting each moves nothing graded
    meta_d = json.load(open(meta_path))["distractor_files"] if meta_path else list(K.DISTRACTORS)
    for gone in meta_d:
        tmp = tempfile.mkdtemp(dir=os.path.dirname(os.path.abspath(target)), suffix=".tmp")
        try:
            for f in os.listdir(target):
                if f != gone:
                    os.symlink(os.path.abspath(os.path.join(target, f)), os.path.join(tmp, f))
            L3 = C.load(tmp)
            ok(graded(C.forecast(L3)) == golden_tuple(), "deleting %s moves the call" % gone)
            ok(C.ask_d(L3) == gold_asks["D sites"] and C.ask_a(L3) == gold_asks["A first lodged"] and
               C.ask_b(L3) == gold_asks["B storage"] and C.ask_c(L3) == gold_asks["C inverter"], "deleting %s moves an ask" % gone)
        finally:
            shutil.rmtree(tmp)

    # ---------------------------------------------------------------- input gates
    ok(len(files) >= 10, "%d files" % len(files))
    ok(len({f.rsplit(".", 1)[1] for f in files}) >= 3, "formats")
    big = max(sum(1 for _ in open(os.path.join(target, f), encoding="utf-8")) - 1 for f in files if f.endswith(".csv"))
    ok(big >= 25000, "largest file %d rows" % big)
    if meta_path:
        meta = json.load(open(meta_path))
        ok(len(meta["distractor_files"]) >= 2 and all(d in files for d in meta["distractor_files"]), "distractors")
    for f in files:
        ok(not re.search(r"distract|decoy|trap|answer|golden|solution", f, re.I), "file name %s" % f)
    man = pd.read_csv(os.path.join(target, K.MANIFEST), dtype=str, keep_default_na=False)
    ok(set(man.owner) - {"Merindah City Council"} <= set(K.PEOPLE.values()), "a manifest owner outside the drawn personas")
    ok(set(man.file) == set(files) - {K.MANIFEST}, "manifest lists %s" % (set(files) ^ set(man.file)))
    dnames = {r[0] for r in wbd["Files"].iter_rows(min_row=2, values_only=True) if r and r[0]}
    ok(dnames == {f for f in files if f.endswith(".csv") and f != K.MANIFEST}, "dictionary files %s" % dnames)
    ext = dict(zip(man.file, man.extracted))
    ok(ext[K.DICTIONARY] >= max(ext[f] for f in dnames), "the dictionary predates an extract it describes")
    import zipfile
    core = zipfile.ZipFile(os.path.join(target, K.DICTIONARY)).read("docProps/core.xml").decode()
    ok(re.search(r"<dcterms:modified[^>]*>(\d{4}-\d\d-\d\d)", core).group(1) == ext[K.DICTIONARY],
       "the dictionary's properties and the manifest give different dates")
    for r_ in man.itertuples():
        if r_.file.endswith(".csv"):
            df = pd.read_csv(os.path.join(target, r_.file), dtype=str, keep_default_na=False)
            ok(int(r_.rows) == len(df), "manifest rows for %s" % r_.file)
            for col in df.columns:
                if col.endswith("_date") or col in ("first_energised", "last_verified"):
                    vals = df[col][df[col] != ""]
                    ok(vals.max() <= r_.extracted, "%s %s runs to %s after its extract %s" % (r_.file, col, vals.max(), r_.extracted))
    ok(der.registered_date.max() <= K.iso(K.BASE) and der.status_date.max() <= K.iso(K.BASE),
       "the DER extract runs past 31 August 2026")

    # ---------------------------------------------------------------- realism sweep
    counts = [c["systems"] for c in camps]
    ok(len(set(counts)) == len(counts), "two campaigns installed the same number of systems")
    ok(len(set(c["letterboxed"] for c in camps)) == len(camps), "identical letterbox counts")
    ok(len(set(c["n365"] for c in camps)) == len(camps), "identical first-application counts")
    hcs = [L["hca"][f]["hc29"] for f in K.FEEDERS]
    ok(len(set(hcs)) == len(hcs), "two feeders share a hosting capacity")
    ok(sum(abs(h * 10 - round(h * 10)) < 1e-6 for h in hcs) <= 2, "round hosting capacities")
    ok("a 0.5 MW lower minimum lowers hosting capacity by 0.3 to 0.6 MW" in hca_txt, "the assessment's sensitivity sentence")
    rows_hc = {m.group(1): tuple(float(m.group(k)) for k in (4, 5, 6, 7)) for m in re.finditer(
        r"(YRB2\d) ([\d.]+) ([\d,]+) ([\d.]+) ([\d.]+) ([\d.]+) ([\d.]+)", hca_txt)}
    ok(set(rows_hc) == set(K.FEEDERS), "the assessment's table rows %s" % sorted(rows_hc))
    steps26 = []
    for f, (m26, m29, h26, h29) in rows_hc.items():
        dm, dh = m26 - m29, h26 - h29
        lo_, hi_ = sorted((0.6 * dm, 1.2 * dm))
        ok(lo_ - 0.0051 <= dh <= hi_ + 0.0051, "%s 2026 case %.2f against minimum demand" % (f, h26))
        steps26.append(dh)
    ok(min(steps26) < 0 < max(steps26) and len({round(x, 2) for x in steps26}) == 8, "the 2026 case steps alike")
    # the assessment's customers (its network model of 30 June 2026) against the 31 August network model: equal because
    # no NMI was energised, abolished or moved between feeders in July or August 2026 (the judge's cosmetic note)
    cust_hc = {m.group(1): int(m.group(3).replace(",", "")) for m in re.finditer(
        r"(YRB2\d) ([\d.]+) ([\d,]+) ([\d.]+) ([\d.]+) ([\d.]+) ([\d.]+)", hca_txt)}
    net_ = L["net"]
    cust30 = net_[net_.first_energised <= "2026-06-30"].groupby("feeder").size().to_dict()
    ok(cust_hc == {f: int(cust30.get(f, 0)) for f in K.FEEDERS}, "the assessment's customers against the 30 June model")
    ok("network model of 30 June 2026" in hca_txt and "issued 24 July 2026" in hca_txt, "the assessment's model date")
    ok(net_.first_energised.max() <= "2026-06-30", "an NMI first energised after the assessment's model date")
    if W is not None:
        jun30 = dt.date(2026, 6, 30)
        ok(max(n["energised"] for n in W.nmis) <= jun30, "an NMI energised in July or August 2026")
        ok(max(n["temp_from"] for n in W.nmis if n.get("temp")) <= jun30, "a temporary supply opened after 30 June 2026")
        ok(max(o["demolished"] for o in W.old_nmis) <= jun30, "an NMI abolished after 30 June 2026")
        ok(all(st <= jun30 for n in W.nmis for st, _ in n["hist"][1:]), "a load transfer after 30 June 2026")
    report(L, camps, pools, rr, se, fc, rungs, band)
    print("checks: %d assertions passed (call %s at %.3f MW, runner-up %s at %.3f, gap %.3f; Bindarra Waters %.3f MW "
          "(%.1f systems), Wattleview %.3f MW (%.1f systems))" % (
              N[0], fc["winner"], ex["YRB27"], fc["runner"], ex[fc["runner"]], fc["gap"], fc["add_mw"]["Bindarra Waters"],
              fc["lift"]["Bindarra Waters"], fc["add_mw"]["Wattleview"], fc["lift"]["Wattleview"]))
    return N[0]


def report(L, camps, pools, rr, se, fc, rungs, band):
    """The campaigns, the ladder and the fork rows as the design note quotes them (printed on every build)."""
    print("campaigns: key | houses without PV at opening | installed | first applications 12 months before | in window "
          "| expected without | added | added per 100 houses")
    for c in camps:
        print("  %s %d | %d | %d | %d | %.1f | %.1f | %.2f" % (c["key"], c["pool"], c["systems"], c["pre"], c["post"],
                                                            c["cf"], c["net"], 100 * c["net"] / c["pool"]))
    for n, p in pools.items():
        print("  %s %d houses without PV | %d letterboxed | %d first applications in the year to August 2026"
              % (n, p["pool"], p["letterboxed"], p["n365"]))
    print("added per house without PV %.5f (se %.5f, cv %.3f); installed per house %.5f" % (
        rr["slope"], se, rr["cv"], C.rate_of(camps, "systems")["slope"]))
    for name, v in rungs.items():
        e = v["excess"]
        print("rung %s: %s %.3f, %s %.3f, YRB27 %.3f (%d), YRB24 %.3f (%d); Bindarra %.3f MW, Wattleview %.3f MW" % (
            name, v["winner"], e[v["winner"]], v["runner"], e[v["runner"]], e["YRB27"], v["order"].index("YRB27") + 1,
            e["YRB24"], v["order"].index("YRB24") + 1, v["add_mw"]["Bindarra Waters"], v["add_mw"]["Wattleview"]))
    print("golden: " + ", ".join("%s %.3f" % (f, fc["excess"][f]) for f in fc["order"]))
    for k, v in SF_SEEN.items():
        print("  %s: %s" % (k, v))


if __name__ == "__main__":
    import sys
    here = os.path.dirname(os.path.abspath(__file__))
    task = os.path.dirname(here)
    t = sys.argv[1] if len(sys.argv) > 1 else os.path.join(task, "target")
    run(t, os.path.join(task, "golden"), meta_path=os.path.join(task, "metadata.json"))
