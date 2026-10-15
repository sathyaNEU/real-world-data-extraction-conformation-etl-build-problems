"""Independent verifier for task136. Reads only the shipped bundle (target/) and the golden workbook, imports nothing
from the generator, and recomputes on its own code path (csv, openpyxl, pypdf; no pandas): the base forecast, what each
of the seven campaigns added against suburbs whose campaign had not yet run (and against never-treated suburbs), the
2026 campaigns, every feeder's excess, the call, the default path, the textbook transfer and each suburb's own before
and after, and the four DER Register figures by feeder.
Exits non-zero on any disagreement with the golden.

    python3 task136/generator/verify.py [target] [golden]"""
from __future__ import annotations

import csv
import datetime as dt
import os
import re
import statistics
import sys
from collections import Counter, defaultdict

import openpyxl
from pypdf import PdfReader

BASE = dt.date(2026, 8, 31)             # NPP-114 3.1: the last day of the month before the forecast is prepared
YEARS = 3                               # 31 August 2026 to the 31 August 2029 design point
FYS = [(dt.date(y, 7, 1), dt.date(y + 1, 6, 30)) for y in (2023, 2024, 2025)]
LAST_FY = FYS[-1]
LIVE = {"Connected", "Removed"}
OPEN_2026 = dt.date(2026, 9, 1)         # the council's letter: sign-ups open 1 September 2026
SUBURBS_2026 = ("Bindarra Waters", "Wattleview")
PORTAL = dt.date(2024, 7, 1)

FAIL = []
PASSED = [0]


def check(cond, msg):
    if cond:
        PASSED[0] += 1
    else:
        FAIL.append(msg)


def one(target, prefix):
    hits = [f for f in os.listdir(target) if f.startswith(prefix)]
    assert len(hits) == 1, (prefix, hits)
    return os.path.join(target, hits[0])


def rows(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def day(s):
    return dt.date.fromisoformat(s) if s else None


def pdf(path):
    return re.sub(r"\s+", " ", " ".join(pg.extract_text() or "" for pg in PdfReader(path).pages))


def half_up(x, unit):
    q = x / unit
    return int(q + 0.5) * unit if q >= 0 else -int(-q + 0.5) * unit


def one_dp(x):
    return half_up(x, 0.1)


# ------------------------------------------------------------------------------------------------------------------
# the bundle

def read(target):
    B = {}
    net = rows(one(target, "nmi_network_model_"))
    B["feeder"] = {r["nmi"]: r["feeder"] for r in net}
    B["tariff"] = {r["nmi"]: r["tariff_class"] for r in net}
    B["locality"] = {r["nmi"]: r["locality"] for r in net}
    B["feeders"] = sorted(set(B["feeder"].values()))
    legacy = {r["legacy_application_no"]: float(r["inverter_kw_ac_each"]) * int(r["inverter_count"])
              for r in rows(one(target, "legacy_inverter_equipment"))}
    apps = []
    for r in rows(one(target, "eg_connection_applications_")):
        if r["inverter_kw_ac"]:
            kw = float(r["inverter_kw_ac"])
        elif r["legacy_application_no"]:
            kw = legacy[r["legacy_application_no"]]
        else:
            assert r["status"] == "Withdrawn", r["application_id"]
            kw = None
        apps.append(dict(nmi=r["nmi"], type=r["application_type"], status=r["status"], received=day(r["received_date"]),
                         connected=day(r["connection_date"]), removed=day(r["removal_date"]), kw=kw,
                         program=r["program_ref"], battery=float(r["battery_kwh"]) if r["battery_kwh"] else 0.0))
    B["apps"] = apps
    B["xref"] = defaultdict(list)
    for r in rows(one(target, "nmi_property_xref_")):
        B["xref"][r["property_no"]].append(r["nmi"])
    B["rates"] = rows(one(target, "merindah_rates_roll_"))
    wb = openpyxl.load_workbook(one(target, "solar_streets_program_summary"), data_only=True)
    head, log = None, []
    for row in wb["Campaigns"].iter_rows(values_only=True):
        if row and row[0] == "Year":
            head = list(row)
        elif head and row and isinstance(row[0], (int, float)) and row[1]:
            log.append(dict(zip(head, row)))
    B["log"] = log
    text = pdf(one(target, "solar_streets_2026_letter"))
    B["letter_text"] = text
    B["letter"] = {}
    for s in SUBURBS_2026:
        m = re.search(s + r"[^.]*?([\d,]+) (?:residential )?addresses", text)
        B["letter"][s] = int(m.group(1).replace(",", ""))
    text = pdf(one(target, "yarrabin_hosting_capacity_assessment"))
    B["hc29"], B["hc26"] = {}, {}
    for m in re.finditer(r"(YRB2\d) [\d.]+ [\d,]+ ([-\d.]+) ([-\d.]+) ([\d.]+) ([\d.]+)", text):
        B["hc26"][m.group(1)], B["hc29"][m.group(1)] = float(m.group(4)), float(m.group(5))
    B["der"] = rows(one(target, "der_register_extract_"))
    B["realloc"] = {r["old_nmi"]: r["new_nmi"] for r in rows(one(target, "nmi_reallocations_"))}
    B["npp"] = pdf(one(target, "NPP-114"))
    return B


# ------------------------------------------------------------------------------------------------------------------
# NPP-114 2.2 and 3.2

def capacity(apps, feeder, when):
    latest = {}
    for a in apps:
        if a["status"] in LIVE and a["connected"] and a["connected"] <= when and not (a["removed"] and a["removed"] <= when):
            k = (a["connected"], 0 if a["type"] == "Alteration" else 1)
            if a["nmi"] not in latest or k > latest[a["nmi"]][0]:
                latest[a["nmi"]] = (k, a["kw"])
    out = defaultdict(float)
    for n, (_, kw) in latest.items():
        if n in feeder:
            out[feeder[n]] += kw
    return out


def base_forecast(B):
    apps, feeder = B["apps"], B["feeder"]
    cap = capacity(apps, feeder, BASE)
    yearly = defaultdict(list)
    for s, e in FYS:
        c1, c0 = capacity(apps, feeder, e), capacity(apps, feeder, s - dt.timedelta(days=1))
        prog = defaultdict(float)
        for a in apps:
            if a["program"] and a["type"] == "New" and a["status"] in LIVE and s <= a["connected"] <= e and a["nmi"] in feeder:
                prog[feeder[a["nmi"]]] += a["kw"]
        for f in B["feeders"]:
            yearly[f].append(c1[f] - c0[f] - prog[f])
    trend = {f: sum(v) / len(v) for f, v in yearly.items()}
    return cap, trend, {f: cap[f] + YEARS * trend[f] for f in B["feeders"]}


def avg_new(B):
    s, e = LAST_FY
    by = defaultdict(list)
    for a in B["apps"]:
        if a["type"] == "New" and a["status"] in LIVE and s <= a["connected"] <= e and a["nmi"] in B["feeder"]:
            by[B["feeder"][a["nmi"]]].append(a["kw"])
    return {f: sum(v) / len(v) for f, v in by.items()}


# ------------------------------------------------------------------------------------------------------------------
# the campaigns: each suburb's owner-occupied houses and the day each first applied

def add_months(d, n):
    y, m = divmod(d.month - 1 + n, 12)
    return dt.date(d.year + y, m + 1, d.day)


def house_first(B):
    """Locality -> list of first-application dates (None for a house without one) over its owner-occupied houses
    (land use 02, rating R1), each house's earliest non-withdrawn application across its NMIs."""
    first = {}
    for a in B["apps"]:
        if a["status"] == "Withdrawn":
            continue
        if a["nmi"] not in first or a["received"] < first[a["nmi"]]:
            first[a["nmi"]] = a["received"]
    out = defaultdict(list)
    for r in B["rates"]:
        if r["land_use_code"] == "02" and r["rating_category"] == "R1":
            ds = [first[n] for n in B["xref"].get(r["property_no"], []) if n in first]
            out[r["locality"]].append(min(ds) if ds else None)
    return out


def between(ds, a, b):
    return sum(1 for d in ds if d is not None and a <= d < b)


def without(ds, t):
    return sum(1 for d in ds if d is None or d >= t)


def campaign_rows(B, HF, mode="nyt", pre=12, prepost=False):
    """For each past campaign: the houses without PV at the opening, the program-referenced systems, the first
    applications in the 12 months before and in the 36 months from the opening (to the base date when sooner), what
    the suburb would have lodged without the campaign (its own pre-year count scaled by the change in suburbs whose
    campaign had not yet run), and what the campaign added."""
    refs = defaultdict(list)
    for a in B["apps"]:
        if a["program"] and a["type"] == "New" and a["status"] in LIVE:
            refs[a["program"]].append(a)
    ref_suburb = {}
    for ref, xs in refs.items():
        locs = Counter(B["locality"][a["nmi"]] for a in xs)
        assert len(locs) == 1, (ref, locs)
        ref_suburb[next(iter(locs))] = ref
    opening = {}
    for row in B["log"]:
        d = row["Sign-ups opened"]
        opening[row["Suburb"]] = d.date() if isinstance(d, dt.datetime) else d
    for s in SUBURBS_2026:
        opening[s] = OPEN_2026
    campaign_suburbs = set(opening)
    never = sorted({loc for loc in B["locality"].values()} - campaign_suburbs)
    out = []
    for row in B["log"]:
        name = row["Suburb"]
        op = opening[name]
        end = min(add_months(op, 36), BASE + dt.timedelta(days=1))
        p0 = add_months(op, -pre)
        ds = HF[name.upper()]
        post, pre_n = between(ds, op, end), between(ds, p0, op)
        if prepost:
            cf = pre_n * ((end - op).days / 365.25) / (pre / 12.0)
        else:
            comp = [s for s, o in opening.items() if o >= end and s != name] if mode in ("nyt", "both") else []
            if mode in ("never", "both"):
                comp = comp + never
            cp = sum(between(HF[s.upper()], op, end) for s in comp)
            cq = sum(between(HF[s.upper()], p0, op) for s in comp)
            cf = pre_n * cp / cq
        res = [n for n, loc in B["locality"].items() if loc == name and B["tariff"][n] == "Residential"]
        out.append(dict(suburb=name, opening=op, pool=without(ds, op), systems=len(refs[ref_suburb[name]]),
                        systems_log=int(row["Systems installed"]), letterboxed=int(row["Homes letterboxed"]),
                        res=len(res), oo=len(ds), post=post, pre=pre_n, cf=cf, net=post - cf))
    return out, never


def pools_2026(B, HF):
    out = {}
    for name in SUBURBS_2026:
        ds = HF[name.upper()]
        res = [n for n, loc in B["locality"].items() if loc == name and B["tariff"][n] == "Residential"]
        out[name] = dict(feeder=suburb_feeder(B, name), pool=without(ds, OPEN_2026), oo=len(ds),
                         letterboxed=B["letter"][name], res=len(res))
    return out


def suburb_feeder(B, suburb):
    c = Counter(B["feeder"][n] for n, loc in B["locality"].items() if loc == suburb)
    assert len(c) == 1, (suburb, c)
    return next(iter(c))


def pooled_lift(rows, pl, measure, basis):
    rate = sum(c[measure] for c in rows) / sum(c[basis] for c in rows)
    return rate, {k: rate * p[basis] for k, p in pl.items()}


def forecast(B, base, avg, pl, lift):
    fc = dict(base)
    add = {}
    for k, n in lift.items():
        fc[pl[k]["feeder"]] += n * avg[pl[k]["feeder"]]
        add[k] = n * avg[pl[k]["feeder"]] / 1000.0
    ex = {f: fc[f] / 1000.0 - B["hc29"][f] for f in B["feeders"]}
    order = sorted(B["feeders"], key=lambda f: -ex[f])
    return fc, ex, order, add


# ------------------------------------------------------------------------------------------------------------------
# the DER Register figures by feeder

def site_key(B, nmi):
    k = nmi.strip()
    if len(k) == 11:
        k = k[:10]
    return B["realloc"].get(k, k)


def der_figures(B):
    """Per feeder: active sites (an NMI with an active AC connection, placed by the network model), the sites whose
    first AC connection of any status was lodged in FY2025-26, the active connections' inverter capacity (a
    two-inverter connection from the older portal carries the connection's total on each row) and their battery
    storage (a Sungrow SBR stack carries the stack's total on each module row)."""
    conns = {}
    for r in B["der"]:
        c = conns.setdefault(r["ac_connection_id"], dict(key=site_key(B, r["nmi"]), status=r["ac_connection_status"],
                                                         reg=day(r["registered_date"]), inv=[], bat=[]))
        if r["device_type"] == "Inverter":
            c["inv"].append((r["model"], float(r["rated_power_kw"])))
        else:
            c["bat"].append((r["model"], float(r["storage_capacity_kwh"])))
    first = {}
    for c in conns.values():
        if c["key"] not in first or c["reg"] < first[c["key"]]:
            first[c["key"]] = c["reg"]
    sites, new, inv, bat = Counter(), Counter(), defaultdict(float), defaultdict(float)
    seen = set()
    for c in conns.values():
        if c["status"] != "Active" or c["key"] not in B["feeder"]:
            continue
        f = B["feeder"][c["key"]]
        if c["key"] not in seen:
            seen.add(c["key"])
            sites[f] += 1
            if LAST_FY[0] <= first[c["key"]] <= LAST_FY[1]:
                new[f] += 1
        rated = [kw for _, kw in c["inv"]]
        m = re.search(r"(\d+\.\d)$", c["inv"][0][0]) if c["inv"] else None
        if len(rated) > 1 and m and abs(rated[0] - float(m.group(1)) * len(rated)) < 0.011:
            inv[f] += rated[0]
        else:
            inv[f] += sum(rated)
        if c["bat"]:
            if c["bat"][0][0] == "SBR HV battery":
                bat[f] += c["bat"][0][1]
            else:
                bat[f] += sum(k for _, k in c["bat"])
    return sites, new, inv, bat


# ------------------------------------------------------------------------------------------------------------------

def golden_book(golden):
    wb = openpyxl.load_workbook(os.path.join(golden, "feeder_upgrade_2029.xlsx"), data_only=True)
    ws = wb["Call"]
    call = [ws.cell(row=i, column=2).value for i in range(4, 11)]
    fcs = {}
    ws = wb["Forecast"]
    for i in range(4, 12):
        fcs[ws.cell(row=i, column=1).value] = [ws.cell(row=i, column=j).value for j in range(2, 11)]
    ws = wb["DER Register"]
    der = {ws.cell(row=i, column=1).value: tuple(ws.cell(row=i, column=j).value for j in range(2, 6)) for i in range(4, 12)}
    return call, fcs, der


def main(target, golden):
    FAIL.clear()
    PASSED[0] = 0
    B = read(target)
    feeders = B["feeders"]
    check(len(feeders) == 8 and set(B["hc29"]) == set(feeders), "eight feeders with a 2029 hosting capacity")
    check("the impact on the feeder's connected capacity at the design point of each program" in B["npp"], "NPP-114 4.1")
    check("rate taken from past programs" not in B["npp"].lower(), "NPP-114 files a pooled rate")
    check("1 September 2026" in B["letter_text"], "the letter's opening date")
    cap, trend, base = base_forecast(B)
    avg = avg_new(B)
    HF = house_first(B)
    camps, never = campaign_rows(B, HF)
    pl = pools_2026(B, HF)
    check(len(camps) == 7 and len(never) >= 5, "seven past campaigns, suburbs never treated %s" % never)
    for c in camps:
        check(c["systems"] == c["systems_log"], "%s: program references %d, council log %d" % (c["suburb"], c["systems"], c["systems_log"]))
        check(c["letterboxed"] == c["res"], "%s: letterboxed %d, residential NMIs %d" % (c["suburb"], c["letterboxed"], c["res"]))
        check(c["systems"] >= 1.7 * c["net"], "%s installed only %.2f times what it added" % (c["suburb"], c["systems"] / c["net"]))
    for k, p in pl.items():
        check(p["letterboxed"] == p["res"], "%s: letter %d, residential NMIs %d" % (k, p["letterboxed"], p["res"]))

    # the call: what each campaign added per house without PV at the opening, pooled, on each 2026 suburb
    rate, lift = pooled_lift(camps, pl, "net", "pool")
    fc, ex, order, add = forecast(B, base, avg, pl, lift)
    w, ru = order[0], order[1]
    gap = ex[w] - ex[ru]
    check(w == "YRB27" and ru == "YRB24", "call %s, runner-up %s" % (w, ru))
    call, fcs, der_g = golden_book(golden)
    check(call[0] == w and call[2] == ru, "golden feeders %s %s" % (call[0], call[2]))
    for got, want, name in ((call[1], ex[w], "excess"), (call[3], ex[ru], "runner-up excess"), (call[4], gap, "gap"),
                            (call[5], add["Bindarra Waters"], "Bindarra Waters MW"),
                            (call[6], add["Wattleview"], "Wattleview MW")):
        check(abs(got - one_dp(want)) < 1e-9, "golden %s %s, verifier %.4f" % (name, got, want))
        k = want * 10
        check(abs((k - int(k)) - 0.5) >= 0.3, "%s %.4f sits near a rounding edge" % (name, want))
    for f in feeders:
        row = fcs[f]
        check(abs(row[0] - cap[f]) <= 0.5 + 1e-6 and abs(row[2] - base[f]) <= 0.5 + 1e-6 and abs(row[5] - fc[f]) <= 0.5 + 1e-6,
              "%s forecast row %s against %.1f %.1f %.1f" % (f, row, cap[f], base[f], fc[f]))
        check(abs(row[7] - one_dp(ex[f])) < 1e-9 and row[8] == order.index(f) + 1, "%s excess and rank" % f)
    graded = (w, ru, one_dp(ex[w]), one_dp(ex[ru]), one_dp(gap), one_dp(add["Bindarra Waters"]), one_dp(add["Wattleview"]))

    def grade(lf):
        _, e2, o2, a2 = forecast(B, base, avg, pl, lf)
        return (o2[0], o2[1], one_dp(e2[o2[0]]), one_dp(e2[o2[1]]), one_dp(e2[o2[0]] - e2[o2[1]]),
                one_dp(a2["Bindarra Waters"]), one_dp(a2["Wattleview"])), e2, o2
    # the comparison read against never-treated suburbs, or both, and other pre-periods: the same figures
    for mode in ("never", "both"):
        for pre in (12, 24, 36):
            rows_, _ = campaign_rows(B, HF, mode=mode, pre=pre)
            g2, _, _ = grade(pooled_lift(rows_, pl, "net", "pool")[1])
            check(g2 == graded, "comparison %s, %d months before, files %s" % (mode, pre, g2))
    for pre in (24, 36):
        rows_, _ = campaign_rows(B, HF, pre=pre)
        g2, _, _ = grade(pooled_lift(rows_, pl, "net", "pool")[1])
        check(g2 == graded, "not yet treated, %d months before, files %s" % (pre, g2))
    # the default path (installed systems per home letterboxed), the textbook transfer (per house without PV) and
    # each suburb's own before and after
    g0_, e0, o0 = grade(pooled_lift(camps, pl, "systems", "letterboxed")[1])
    check(o0[0] == "YRB24" and e0["YRB24"] - e0["YRB27"] >= 0.1, "default path files %s" % o0[0])
    g1_, e1, o1 = grade(pooled_lift(camps, pl, "systems", "pool")[1])
    check(o1[0] == "YRB24" and e1["YRB24"] - e1["YRB27"] >= 0.1, "the textbook transfer files %s" % o1[0])
    for pre in (12, 24, 36):
        rows_, _ = campaign_rows(B, HF, pre=pre, prepost=True)
        g3, e3, o3 = grade(pooled_lift(rows_, pl, "net", "pool")[1])
        check(o3[0] == "YRB24", "before and after (%d months) files %s" % (pre, o3[0]))
    g4, _, o4 = grade({k: 0.0 for k in pl})
    check(o4[0] not in ("YRB27", "YRB24"), "no campaigns files %s" % o4[0])
    # NPP-114 5.1 and the prompt name the design year's case; the 2026 case beside it is another quantity
    check("for the design year's spring minimum-demand case" in B["npp"], "NPP-114 5.1")
    _, e5, o5, _ = forecast(dict(B, hc29=B["hc26"]), base, avg, pl, lift)
    check(o5[0] != w or abs(e5[w] - ex[w]) > 0.05, "the 2026 case gives the 2029 figures")

    # the DER Register figures
    sites, new, inv, bat = der_figures(B)
    for f in feeders:
        want = (sites[f], new[f], half_up(inv[f], 10), half_up(bat[f], 10))
        check(tuple(der_g[f]) == want, "%s DER Register row golden %s, verifier %s" % (f, der_g[f], want))
        for x, name in ((inv[f], "inverter"), (bat[f], "storage")):
            k = x / 10
            check(abs((k - int(k)) - 0.5) >= 0.07, "%s %s %.2f near a rounding edge" % (f, name, x))
    check(os.path.getsize(os.path.join(golden, "yarrabin_pv_vs_hosting_2029.png")) > 20000, "the chart")

    if FAIL:
        for m in FAIL:
            print("VERIFY FAIL:", m)
        raise SystemExit(1)
    print("verify: %d checks passed from the bundle alone (%s %.3f MW above hosting capacity, runner-up %s %.3f, gap "
          "%.3f; Bindarra Waters %.3f MW and Wattleview %.3f MW; %.5f systems added per house without PV)"
          % (PASSED[0], w, ex[w], ru, ex[ru], gap, add["Bindarra Waters"], add["Wattleview"], rate))
    return PASSED[0]


if __name__ == "__main__":
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    t = sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "target")
    g = sys.argv[2] if len(sys.argv) > 2 else os.path.join(here, "golden")
    main(t, g)
