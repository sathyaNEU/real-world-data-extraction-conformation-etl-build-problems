"""The pack's documents: NPP-114 (the distributor's planning procedure), the 2026 hosting capacity assessment, the
council's letter on the 2027 campaigns and its program workbook, the data dictionary and the network's peak demand
forecast."""
from __future__ import annotations

import datetime as dt

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import A4                                                # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import mm                                                    # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import common as K   # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/"
for _n, _f in (("Carlito", "crosextra/Carlito-Regular.ttf"), ("Carlito-Bold", "crosextra/Carlito-Bold.ttf"),
               ("Carlito-Italic", "crosextra/Carlito-Italic.ttf"),
               ("LiberationSerif", "liberation/LiberationSerif-Regular.ttf"),
               ("LiberationSerif-Bold", "liberation/LiberationSerif-Bold.ttf")):
    pdfmetrics.registerFont(TTFont(_n, FONT_DIR + _f))
SANS, SANS_B, SANS_I = "Carlito", "Carlito-Bold", "Carlito-Italic"
SERIF, SERIF_B = "LiberationSerif", "LiberationSerif-Bold"
INK = colors.HexColor("#1d2329")
GREY = colors.HexColor("#59626b")
TEAL = colors.HexColor("#0b5d6b")
RULE = colors.HexColor("#9aa4ad")

BODY = ParagraphStyle("body", fontName=SANS, fontSize=10.2, leading=13.4, textColor=INK, spaceAfter=5)
CLAUSE = ParagraphStyle("clause", parent=BODY, leftIndent=26, firstLineIndent=-26)
H1 = ParagraphStyle("h1", fontName=SANS_B, fontSize=15, leading=19, textColor=TEAL, spaceAfter=2)
H1S = ParagraphStyle("h1s", fontName=SANS, fontSize=10.5, leading=13.5, textColor=GREY, spaceAfter=10)
H2 = ParagraphStyle("h2", fontName=SANS_B, fontSize=11, leading=14, textColor=INK, spaceBefore=7, spaceAfter=3)
NOTE = ParagraphStyle("note", fontName=SANS_I, fontSize=8.8, leading=11.2, textColor=GREY, spaceBefore=3)
CELL = ParagraphStyle("cell", fontName=SANS, fontSize=8.6, leading=10.4, textColor=INK)
CELLB = ParagraphStyle("cellb", parent=CELL, fontName=SANS_B)
LTR = ParagraphStyle("ltr", fontName=SERIF, fontSize=10.6, leading=14.2, textColor=INK, spaceAfter=7)
LTRB = ParagraphStyle("ltrb", parent=LTR, fontName=SERIF_B)


def _footer(left, right):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont(SANS, 7.8)
        canvas.setFillColor(GREY)
        canvas.drawString(20 * mm, 11 * mm, left)
        canvas.drawRightString(A4[0] - 20 * mm, 11 * mm, "%s  |  Page %d" % (right, doc.page))
        canvas.restoreState()
    return draw


def _doc(path, title, author):
    return SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                             bottomMargin=20 * mm, title=title, author=author, subject="", creator="", producer="")


def _table(rows, widths, head=1, align_right_from=None):
    data = [[Paragraph(str(c), CELLB if i < head else CELL) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, repeatRows=head)
    style = [("LINEBELOW", (0, head - 1), (-1, head - 1), 0.6, RULE), ("LINEBELOW", (0, -1), (-1, -1), 0.4, RULE),
             ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 1.8),
             ("BOTTOMPADDING", (0, 0), (-1, -1), 1.8), ("LEFTPADDING", (0, 0), (-1, -1), 3)]
    t.setStyle(TableStyle(style))
    return t


# ------------------------------------------------------------------------------------------------------------------
# NPP-114

def write_npp(path):
    s = [Paragraph("Mirrabee Energy", H1S),
         Paragraph("NPP-114 Forecasting micro embedded generation for distribution augmentation", H1),
         Paragraph("Network Planning Procedure. Revision 3, effective 2 March 2026. Owner: Manager Network Planning "
                   "(%s). Supersedes revision 2 (April 2024)." % K.PEOPLE["planning"], H1S)]
    s.append(Paragraph("1 Purpose and scope", H2))
    s.append(Paragraph("1.1&nbsp;&nbsp;This procedure sets how Network Planning forecasts the micro embedded generation "
                       "connected to an 11 kV or 22 kV feeder at the design point of an augmentation, and how a "
                       "hosting-capacity augmentation is assigned when the capital works program funds fewer "
                       "augmentations than there are feeders that need one.", CLAUSE))
    s.append(Paragraph("1.2&nbsp;&nbsp;It does not apply to large embedded generation (above 30 kVA), which the hosting "
                       "capacity assessment models at its export limit.", CLAUSE))
    s.append(Paragraph("2 Definitions", H2))
    for no, text in [
        ("2.1", "<b>Micro EG</b>: an inverter-connected generating system of 30 kVA or less connected to the "
                "distribution network under a micro embedded generation connection agreement, at a residential or "
                "business NMI."),
        ("2.2", "<b>Connected capacity</b> of a feeder at a date: the sum, over the NMIs the network model places on the "
                "feeder at the base date, of the inverter capacity (kW, AC rating) of each NMI's micro EG as last "
                "connected on or before that date. An alteration replaces the configuration it alters. A system "
                "removed on or before the date is not counted."),
        ("2.3", "<b>Program system</b>: a micro EG system installed under a council, retailer or government program "
                "and lodged with the program's reference."),
        ("2.4", "<b>Design point</b>: 31 August of the design year, the start of the spring minimum-demand season."),
        ("2.5", "<b>Financial year</b>: 1 July to 30 June."),
    ]:
        s.append(Paragraph("%s&nbsp;&nbsp;%s" % (no, text), CLAUSE))
    s.append(Paragraph("3 Base forecast", H2))
    s.append(Paragraph("3.1&nbsp;&nbsp;The base date is the last day of the month before the forecast is prepared. "
                       "The network model at the base date places each NMI on its feeder.", CLAUSE))
    s.append(Paragraph("3.2&nbsp;&nbsp;The base forecast at the design point is the connected capacity at the base "
                       "date plus, for each year from the base date to the design point, the feeder's trend: the "
                       "average, over the last three complete financial years, of the year's increase in connected "
                       "capacity less the capacity of the program systems first connected in that year.", CLAUSE))
    s.append(Paragraph("4 Programs", H2))
    s.append(Paragraph("4.1&nbsp;&nbsp;A program scheduled on a feeder, whose sign-up window closes at least six months "
                       "before the design point, adds to the base forecast the number of systems that the "
                       "distributor's records of past programs of the same kind indicate it will bring.", CLAUSE))
    s.append(Paragraph("4.2&nbsp;&nbsp;Each such system is taken at the average inverter capacity of the micro EG "
                       "systems first connected on that feeder in the last complete financial year.", CLAUSE))
    s.append(Paragraph("4.3&nbsp;&nbsp;A program that has not been confirmed in writing by its sponsor is not "
                       "forecast.", CLAUSE))
    s.append(Paragraph("5 Hosting capacity", H2))
    s.append(Paragraph("5.1&nbsp;&nbsp;A feeder's hosting capacity is the figure in the current hosting capacity "
                       "assessment for the design year's spring minimum-demand case.", CLAUSE))
    s.append(Paragraph("6 Assignment", H2))
    s.append(Paragraph("6.1&nbsp;&nbsp;Where the capital works program funds one hosting-capacity augmentation, it is "
                       "assigned to the feeder whose forecast connected capacity at the design point exceeds its "
                       "hosting capacity by the most.", CLAUSE))
    s.append(Paragraph("7 Reporting", H2))
    s.append(Paragraph("7.1&nbsp;&nbsp;Forecasts are carried unrounded and rounded once for reporting: capacity and "
                       "excess in MW to one decimal, systems to the nearest ten.", CLAUSE))
    s.append(Spacer(1, 8))
    s.append(Paragraph("Revision 3 moved the design point from 1 October to 31 August and added clause 4.3. Approved: "
                       "%s, Manager Network Planning, 26 February 2026." % K.PEOPLE["planning"], NOTE))
    _doc(path, "NPP-114 Forecasting micro embedded generation for distribution augmentation (rev 3)",
         K.PEOPLE["planning"]).build(s, onFirstPage=_footer("Mirrabee Energy  |  NPP-114 rev 3", "Uncontrolled when printed"),
                                     onLaterPages=_footer("Mirrabee Energy  |  NPP-114 rev 3", "Uncontrolled when printed"))


# ------------------------------------------------------------------------------------------------------------------
# the hosting capacity assessment

def write_hca(path, table):
    """table: feeder -> dict(route_km, customers, min26, min29, hc26, hc29)."""
    s = [Paragraph("Buchanan Power Engineering", H1S),
         Paragraph("Yarrabin Zone Substation 22 kV Feeders: Hosting Capacity Assessment 2026", H1),
         Paragraph("Prepared for Mirrabee Energy, Network Planning. Report BPE-2026-031, issued 24 July 2026.", H1S)]
    s.append(Paragraph("1 Scope", H2))
    s.append(Paragraph("This assessment gives, for each of the eight 22 kV feeders supplied from Yarrabin zone "
                       "substation, the capacity of micro embedded generation (inverter systems of 30 kVA or less) "
                       "the feeder can host without a voltage or thermal limit being exceeded, for the spring "
                       "minimum-demand case of 2026 and of the 2029 design year.", BODY))
    s.append(Paragraph("2 Method", H2))
    for t in ["Load flow on each feeder's network model (Mirrabee GIS extract of 30 June 2026), with the spring "
              "minimum daytime demand set from 2025 feeder-head SCADA and the network's demand forecast for 2029.",
              "Steady-state voltage limit of 253 V at the customer's point of supply (AS 61000.3.100); conductor and "
              "distribution transformer thermal ratings from the asset register.",
              "Micro EG modelled with volt-var and volt-watt responses to AS/NZS 4777.2:2020 and spread across each "
              "feeder's distribution transformers in proportion to customers.",
              "Large embedded generation connected or committed at 30 June 2026 modelled at its export limit; the "
              "hosting capacity below is what remains for micro EG.",
              "No network augmentation after 30 June 2026 is assumed."]:
        s.append(Paragraph("&bull;&nbsp;&nbsp;" + t, CLAUSE))
    s.append(Paragraph("3 Results", H2))
    rows = [["Feeder", "Route length (km)", "Customers", "Spring minimum daytime demand 2026 (MW)",
             "Spring minimum daytime demand 2029 (MW)", "Hosting capacity, 2026 case (MW)",
             "Hosting capacity, 2029 case (MW)"]]
    for f in K.FEEDERS:
        t = table[f]
        rows.append([f, "%.1f" % t["route_km"], "{:,}".format(t["customers"]), "%.2f" % t["min26"],
                     "%.2f" % t["min29"], "%.2f" % t["hc26"], "%.2f" % t["hc29"]])
    s.append(_table(rows, [16 * mm, 20 * mm, 20 * mm, 30 * mm, 30 * mm, 27 * mm, 27 * mm]))
    s.append(Paragraph("Customers are NMIs supplied from the feeder in the network model of 30 June 2026. Hosting "
                       "capacity is stated for micro EG and does not depend on how much is connected today.", NOTE))
    s.append(Paragraph("4 Limits", H2))
    s.append(Paragraph("On every feeder the binding limit in the 2029 case is voltage rise on the low-voltage network "
                       "at the far end of the feeder during the spring minimum. Results are sensitive to the "
                       "minimum-demand forecast: a 0.5 MW lower minimum lowers hosting capacity by 0.3 to 0.6 MW, "
                       "depending on the feeder.", BODY))
    s.append(Paragraph("Prepared by %s, Principal Engineer, Buchanan Power Engineering." % K.PEOPLE["consultant"], NOTE))
    _doc(path, "Yarrabin Zone Substation 22 kV Feeders: Hosting Capacity Assessment 2026", K.PEOPLE["consultant"]).build(
        s, onFirstPage=_footer("BPE-2026-031", "Buchanan Power Engineering"),
        onLaterPages=_footer("BPE-2026-031", "Buchanan Power Engineering"))


# ------------------------------------------------------------------------------------------------------------------
# the council's letter

def write_letter(path, counts):
    s = [Paragraph("<b>Merindah City Council</b>", ParagraphStyle("lh", fontName=SERIF_B, fontSize=16, leading=20,
                                                                  textColor=TEAL)),
         Paragraph("Sustainability Programs, PO Box 412, Merindah QLD 4510", ParagraphStyle(
             "lh2", fontName=SERIF, fontSize=9, leading=12, textColor=GREY, spaceAfter=16)),
         Paragraph("14 August 2026", LTR),
         Paragraph("Network Planning<br/>Mirrabee Energy<br/>Locked Bag 7<br/>Merindah QLD 4510", LTR),
         Paragraph("Our reference: SS/2027/01", LTR),
         Paragraph("Solar Streets 2027", LTRB),
         Paragraph("Council's 2026-27 budget, adopted on 30 July 2026, includes two Solar Streets bulk-buy campaigns in "
                   "spring 2027. Sign-ups will open on Monday 6 September 2027 and close on Friday 26 November 2027 in "
                   "both suburbs:", LTR),
         Paragraph("&bull;&nbsp;&nbsp;Bindarra Waters, where the letterbox drop will go to all {:,} residential "
                   "addresses;".format(counts["Bindarra Waters"]), LTR),
         Paragraph("&bull;&nbsp;&nbsp;Wattleview, where the letterbox drop will go to all {:,} residential "
                   "addresses.".format(counts["Wattleview"]), LTR),
         Paragraph("Council will appoint the installer by tender in February 2027. As in earlier campaigns, the "
                   "installer will lodge each connection application with the program reference, and installations "
                   "will follow over the summer after sign-ups close. Information sessions will be held at the "
                   "Bindarra Waters community centre and the Wattleview library in early September 2027.", LTR),
         Paragraph("Please treat this letter as Council's written confirmation of the 2027 campaigns for your network "
                   "planning. A summary of the earlier campaigns is attached.", LTR),
         Spacer(1, 10),
         Paragraph("Yours sincerely,", LTR),
         Spacer(1, 16),
         Paragraph("%s<br/>Sustainability Programs Officer" % K.PEOPLE["council"], LTR)]
    _doc(path, "Solar Streets 2027", K.PEOPLE["council"]).build(s)


# ------------------------------------------------------------------------------------------------------------------
# workbooks

HEAD_FILL = PatternFill("solid", fgColor="DDE7EA")


def write_program_log(path, rows):
    """rows: dicts with the summary columns."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Campaigns"
    ws["A1"] = "Merindah City Council"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Solar Streets bulk-buy program: campaign summary 2018 to 2024"
    ws["A2"].font = Font(bold=True, size=11)
    ws["A3"] = "Prepared by Sustainability Programs. Last updated 3 March 2025."
    ws["A3"].font = Font(italic=True, size=9)
    head = ["Year", "Campaign", "Suburb", "Sign-ups opened", "Sign-ups closed", "Appointed installer",
            "Homes letterboxed", "Information sessions", "Registrations of interest", "Site assessments",
            "Contracts signed", "Systems installed", "Panel capacity installed (kW)", "Notes"]
    for j, h in enumerate(head, 1):
        c = ws.cell(row=5, column=j, value=h)
        c.font = Font(bold=True, size=10)
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, r in enumerate(rows, 6):
        vals = [r["year"], r["campaign"], r["suburb"], r["open"], r["close"], r["installer"], r["letterboxed"],
                r["sessions"], r["registrations"], r["assessments"], r["contracts"], r["systems"], r["kwp"], r["notes"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=None if v == "" else v)
            if isinstance(v, dt.date):
                c.number_format = "d mmm yyyy"
            elif j == 13:
                c.number_format = "#,##0.0"
            elif isinstance(v, int) and j >= 7:
                c.number_format = "#,##0"
    last = 5 + len(rows)
    notes = ["1 Homes letterboxed: every residential delivery point in the suburb on Australia Post's list at the drop.",
             "2 Systems installed: as reported by the appointed installer at program close and matched to Mirrabee "
             "Energy's connection records (program reference SS-yy-xxx).",
             "3 Contracts signed include contracts the household later cancelled.",
             "4 No campaign was run in 2020."]
    for k, t in enumerate(notes):
        ws.cell(row=last + 2 + k, column=1, value=t).font = Font(italic=True, size=9)
    widths = [7, 22, 18, 14, 14, 28, 11, 11, 13, 11, 10, 10, 13, 34]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[5].height = 42
    ws.freeze_panes = "A6"
    t = wb.create_sheet("Program terms")
    t["A1"] = "Solar Streets: program terms (as published with each campaign)"
    t["A1"].font = Font(bold=True, size=11)
    terms = [
        "Solar Streets offers owner-occupiers of houses in the campaign suburb a rooftop solar system from Council's "
        "appointed installer at a price negotiated for the campaign.",
        "Council letterboxes every residential address in the suburb, holds information sessions, and keeps sign-ups "
        "open for twelve weeks.",
        "The installer assesses each site, offers a fixed-price contract and installs over the months after sign-ups "
        "close. The household may cancel a contract up to ten business days before installation.",
        "Tenants, and owners of units or townhouses under a body corporate, are referred to the installer's standard "
        "offer outside the campaign.",
        "Council does not receive any payment from the installer and does not own any system installed.",
    ]
    for k, line in enumerate(terms, 3):
        c = t.cell(row=k, column=1, value=line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    t.column_dimensions["A"].width = 110
    wb.save(path)


def write_dictionary(path, counts):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Files"
    ws.append(["File", "Grain", "Rows"])
    grains = [
        (K.REGISTER, "One row per micro EG connection application (a new system or an alteration) at an NMI."),
        (K.LEGACY, "One row per application migrated from the pre-July-2019 connections system: its inverters."),
        (K.NETWORK, "One row per NMI on a Yarrabin feeder in the network model at 31 August 2026."),
        (K.XREF, "One row per NMI: the council property it supplies."),
        (K.RATES, "One row per rateable property in Merindah's western rating area, 2026-27 roll."),
        (K.DER, "One row per device (inverter or battery) of each AC connection the DER Register holds at an NMI "
                "on a Yarrabin feeder."),
        (K.REALLOC, "One row per NMI replaced by another NMI at the same premises."),
        (K.REBATE, "One row per home battery rebate approval for a Merindah address in FY2025-26, as the state "
                   "program sent it."),
    ]
    for f, g in grains:
        ws.append([f, g, counts.get(f, "")])
    fields = {
        "Connections register": [
            ("application_id", "Connection application number."),
            ("nmi", "NMI the application is for (the NMI of the premises at the base date)."),
            ("application_type", "New (a first system at the NMI) or Alteration (a change to the system at the NMI)."),
            ("received_date", "Date the application was received."),
            ("approved_date", "Date the application was approved; blank if not yet approved or withdrawn before approval."),
            ("status", "Status at the extract. Under assessment: received, not yet approved. Approved: approved, not "
                       "yet connected. Connected: connected (connection_date). Withdrawn: withdrawn by the applicant "
                       "or lapsed before connection; nothing was connected under the application. Removed: connected, "
                       "then removed (removal_date)."),
            ("status_date", "Date of the status shown."),
            ("connection_date", "Date the system was connected (meter reconfigured); blank until connected."),
            ("removal_date", "Date a connected system was removed; blank otherwise."),
            ("inverter_kw_ac", "Total inverter capacity (kW, AC rating) of the system at the NMI after the work in this "
                               "application. Blank on applications migrated from the legacy system."),
            ("panel_kwp", "Total panel capacity (kWp) after the work in this application."),
            ("battery_kwh", "Battery storage capacity (kWh) after the work in this application, if any."),
            ("export_limit_kw", "Export limit in the connection agreement (applications from 2019)."),
            ("phases", "Number of phases of the inverter connection."),
            ("feeder_at_application", "Feeder recorded on the application when it was received."),
            ("installer_accreditation_no", "Installer's accreditation number."),
            ("program_ref", "Program reference lodged with the application (SS-yy-xxx: Merindah Solar Streets)."),
            ("legacy_application_no", "Application number in the legacy system, for applications migrated in June 2019."),
        ],
        "Network model": [
            ("nmi", "National Metering Identifier (ten characters)."), ("feeder", "Feeder supplying the NMI at 31 August 2026."),
            ("zone_substation", "Zone substation."), ("distribution_transformer", "Distribution transformer supplying the NMI."),
            ("tariff_class", "Residential or Business network tariff class."),
            ("street_address", "Supply address."), ("locality", "Locality of the supply address."), ("postcode", "Postcode."),
            ("nmi_status", "Active or De-energised."), ("first_energised", "Date the NMI was first energised."),
        ],
        "Cross-reference": [
            ("nmi", "NMI."), ("property_no", "Council property number the NMI supplies."),
            ("match_method", "How GIS matched the NMI to the property: Address, Spatial or Manual."),
            ("last_verified", "Date the match was last verified."),
        ],
        "Rates roll": [
            ("property_no", "Council property number."), ("lot_on_plan", "Lot on plan."),
            ("property_address", "Property address."), ("locality", "Locality (upper case on the roll)."),
            ("postcode", "Postcode."), ("land_use_code", "Valuer-General primary land use code."),
            ("land_use_description", "Land use description."),
            ("rating_category", "Council differential general rating category for 2026-27."),
            ("rating_category_description", "Rating category description."),
            ("land_area_m2", "Land area (square metres); blank for building units."),
            ("dwelling_year", "Year the dwelling was completed, from building approvals; blank where none."),
        ],
        "DER Register extract": [
            ("ac_connection_id", "DER Register AC connection identifier."),
            ("nmi", "NMI recorded for the AC connection by the installer."),
            ("device_type", "Inverter or Battery."), ("manufacturer", "Device manufacturer."), ("model", "Device model."),
            ("rated_power_kw", "Inverter rated AC power (kW), inverter rows."),
            ("storage_capacity_kwh", "Storage capacity (kWh), battery rows."),
            ("ac_connection_status", "Active or Decommissioned."), ("status_date", "Date of the status shown."),
            ("registered_date", "Date the installer lodged the AC connection."),
            ("installer_accreditation_no", "Installer's accreditation number."),
            ("feeder", "Feeder, added by Mirrabee's extraction from its NMI history."),
        ],
        "NMI re-allocations": [
            ("old_nmi", "NMI no longer in use."), ("new_nmi", "NMI that replaced it at the same premises."),
            ("effective_date", "Date of the replacement."), ("reason", "Reason recorded."),
        ],
        "Battery rebate approvals": [
            ("approval_ref", "State program approval reference."), ("property_address", "Address on the application."),
            ("locality", "Locality."), ("postcode", "Postcode."),
            ("battery_make_model", "Battery quoted on the application."),
            ("nominal_kwh", "Nominal storage capacity of the quoted battery (kWh)."),
            ("approved_date", "Date the rebate was approved."), ("rebate_aud", "Rebate approved (AUD)."),
            ("approval_status", "Status at the program's export: Approved, Paid or Expired. The rebate is paid "
                                "after installation."),
        ],
    }
    for sheet, rows in fields.items():
        w = wb.create_sheet(sheet)
        w.append(["Field", "Definition"])
        for r in rows:
            w.append(list(r))
        w.column_dimensions["A"].width = 28
        w.column_dimensions["B"].width = 100
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 100
    ws.column_dimensions["C"].width = 10
    wb.save(path)


def write_peak(path, peaks):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summer peak POE50"
    ws["A1"] = "Mirrabee Energy  Feeder demand forecast 2027 to 2031 (summer peak, 50% POE, MVA)"
    ws["A1"].font = Font(bold=True, size=11)
    ws["A2"] = "Network Planning. Forecast of 12 May 2026 for the Distribution Annual Planning Report."
    ws["A2"].font = Font(italic=True, size=9)
    ws.append([])
    ws.append(["Feeder", "Rating (MVA)", "2026-27", "2027-28", "2028-29", "2029-30", "2030-31"])
    for f in K.FEEDERS:
        ws.append([f] + peaks[f])
    for row in ws.iter_rows(min_row=5, min_col=2, max_col=7):
        for c in row:
            c.number_format = "0.00"
    ws.column_dimensions["A"].width = 10
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 11
    wb.save(path)
