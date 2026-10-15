"""The golden deliverables: feeder_upgrade_2029.xlsx (the call, the forecast by feeder, the Solar Streets estimate, the
DER Register figures by feeder) and yarrabin_pv_vs_hosting_2029.png, plus generator/figures.json, every figure computed
from the shipped pack (calc.py)."""
from __future__ import annotations

import json
import os

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import calc as C
import common as K
import postprocess as PP

def r1(x):
    return round(x + 1e-12, 1)


def ten(x):
    return int(round(x / 10.0) * 10)


def compute(target):
    L = C.load(target)
    base = C.base_forecast(L["reg"])
    camps = C.campaigns(L)
    pools = C.pools_2026(L)
    fc = C.forecast(L, base=base, camps=camps, pools=pools)
    default = C.forecast(L, measure="systems", basis="letterboxed", base=base, camps=camps, pools=pools)
    reflex = C.forecast(L, measure="systems", basis="pool", base=base, camps=camps, pools=pools)
    prepost = C.forecast(L, base=base, camps=C.campaigns(L, form="prepost"), pools=pools)
    ratios = {}
    ops = C.openings(L)
    for x in camps:
        op = pd.Timestamp(x["opening"])
        end = C.window_end(op)
        comp = C.comparison(L, op, end, "nyt", x["suburb"])
        cp = sum(C.n_between(C.suburb_first(L, s, "received_date"), op, end) for s in comp)
        cq = sum(C.n_between(C.suburb_first(L, s, "received_date"), op - pd.DateOffset(months=12), op) for s in comp)
        ratios[x["key"]] = dict(comp=comp, post=cp, pre=cq, end=(end - pd.Timedelta(days=1)).date())
    asks = dict(sites=C.ask_d(L), first=C.ask_a(L), storage=C.ask_b(L), inverter=C.ask_c(L))
    return dict(L=L, base=base, camps=camps, pools=pools, fc=fc, default=default, reflex=reflex, prepost=prepost,
                rate=C.rate_of(camps), gross=C.rate_of(camps, "systems"), se=C.ratio_se(camps), ratios=ratios,
                asks=asks, openings=ops)


def path_block(v):
    return {"feeder": v["winner"], "excess_mw": "%.1f" % r1(v["excess"][v["winner"]]), "runner_up": v["runner"],
            "bindarra_waters_mw": "%.1f" % r1(v["add_mw"]["Bindarra Waters"]),
            "wattleview_mw": "%.1f" % r1(v["add_mw"]["Wattleview"])}


def figures(c):
    fc, pools = c["fc"], c["pools"]
    w, ru = fc["winner"], fc["runner"]
    a = c["asks"]
    out = {
        "call": {"feeder": w, "excess_mw": "%.1f" % r1(fc["excess"][w]), "runner_up": ru,
                 "runner_up_excess_mw": "%.1f" % r1(fc["excess"][ru]), "gap_mw": "%.1f" % r1(fc["gap"]),
                 "bindarra_waters_mw": "%.1f" % r1(fc["add_mw"]["Bindarra Waters"]),
                 "wattleview_mw": "%.1f" % r1(fc["add_mw"]["Wattleview"])},
        "steps": {
            "connected_kw": {f: "{:,}".format(int(round(c["base"]["cap"].get(f, 0)))) for f in K.FEEDERS},
            "trend_kw": {f: "{:,}".format(int(round(c["base"]["trend"][f]))) for f in K.FEEDERS},
            "base_kw": {f: "{:,}".format(int(round(c["base"]["base"][f]))) for f in K.FEEDERS},
            "forecast_kw": {f: "{:,}".format(int(round(fc["fc"][f]))) for f in K.FEEDERS},
            "hosting_mw": {f: "%.2f" % c["L"]["hca"][f]["hc29"] for f in K.FEEDERS},
            "excess_mw": {f: "%.1f" % r1(fc["excess"][f]) for f in K.FEEDERS},
            "campaigns": {x["key"]: {"houses_without_pv": "{:,}".format(x["pool"]), "systems": x["systems"],
                                     "first_applications_12_months_before": x["pre"],
                                     "first_applications_in_window": x["post"],
                                     "expected_without_campaign": "%.1f" % x["cf"], "added": "%.1f" % x["net"],
                                     "added_per_100_houses": "%.2f" % (100.0 * x["net"] / x["pool"])}
                          for x in c["camps"]},
            "campaign_systems_total": "{:,}".format(sum(x["systems"] for x in c["camps"])),
            "campaign_houses_without_pv_total": "{:,}".format(sum(x["pool"] for x in c["camps"])),
            "campaign_added_total": "%.1f" % sum(x["net"] for x in c["camps"]),
            "added_per_100_houses": "%.2f" % (100.0 * c["rate"]["slope"]),
            "installed_per_100_houses": "%.2f" % (100.0 * c["gross"]["slope"]),
            "bindarra_houses_without_pv": "{:,}".format(pools["Bindarra Waters"]["pool"]),
            "wattleview_houses_without_pv": "{:,}".format(pools["Wattleview"]["pool"]),
            "bindarra_systems_unrounded": "%.1f" % fc["lift"]["Bindarra Waters"],
            "wattleview_systems_unrounded": "%.1f" % fc["lift"]["Wattleview"],
            "avg_kw_yrb24": "%.2f" % fc["avgkw"]["YRB24"], "avg_kw_yrb27": "%.2f" % fc["avgkw"]["YRB27"],
            "program_kw_yrb24": "{:,}".format(int(round(fc["add_kw"]["YRB24"]))),
            "program_kw_yrb27": "{:,}".format(int(round(fc["add_kw"]["YRB27"]))),
            "winner_forecast_mw": "%.1f" % r1(fc["fc"][w] / 1000.0),
            "winner_hosting_mw": "%.2f" % c["L"]["hca"][w]["hc29"],
        },
        "default_path": path_block(c["default"]),
        "reflex_path": path_block(c["reflex"]),
        "prepost_path": path_block(c["prepost"]),
        "der": {f: {"active_sites": a["sites"][f], "first_lodged_fy2025_26": a["first"][f],
                    "inverter_kw": "{:,}".format(ten(a["inverter"][f])),
                    "storage_kwh": "{:,}".format(ten(a["storage"][f]))} for f in K.FEEDERS},
    }
    return out


HEAD = PatternFill("solid", fgColor="E4ECEF")
CALL = PatternFill("solid", fgColor="F3F0E6")


def _head(ws, row, values, widths=None, height=44):
    for j, v in enumerate(values, 1):
        c = ws.cell(row=row, column=j, value=v if v else None)
        c.font = Font(bold=True, size=10)
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = height
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w


def write_book(path, c, figs):
    fc, pools, L = c["fc"], c["pools"], c["L"]
    w, ru = fc["winner"], fc["runner"]
    st = figs["steps"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Call"
    ws["A1"] = ("%s takes the FY2027-28 hosting-capacity augmentation: its rooftop PV sits furthest above hosting "
                "capacity at the 31 August 2029 design point" % w)
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Yarrabin 22 kV feeders, forecast under NPP-114 rev 3 from a base date of 31 August 2026."
    ws["A2"].font = Font(italic=True, size=9)
    rows = [
        ("Feeder to receive the augmentation", w, None),
        ("Rooftop PV above hosting capacity at 31 August 2029 (MW)", float(figs["call"]["excess_mw"]), "0.0"),
        ("Runner-up feeder", ru, None),
        ("Runner-up above hosting capacity (MW)", float(figs["call"]["runner_up_excess_mw"]), "0.0"),
        ("Gap between them (MW)", float(figs["call"]["gap_mw"]), "0.0"),
        ("Solar Streets 2026, Bindarra Waters (YRB24): PV added by 31 Aug 2029 (MW)",
         float(figs["call"]["bindarra_waters_mw"]), "0.0"),
        ("Solar Streets 2026, Wattleview (YRB27): PV added by 31 Aug 2029 (MW)", float(figs["call"]["wattleview_mw"]),
         "0.0"),
    ]
    for i, (k, v, fmt) in enumerate(rows, 4):
        ws.cell(row=i, column=1, value=k)
        cc = ws.cell(row=i, column=2, value=v)
        cc.fill = CALL
        cc.alignment = Alignment(horizontal="right")
        if fmt:
            cc.number_format = fmt
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 14
    notes = [
        "%s forecast %s MW at 31 Aug 2029 against 2029 hosting capacity %s MW; %s MW of it is the Wattleview "
        "bulk-buy." % (w, st["winner_forecast_mw"], st["winner_hosting_mw"], figs["call"]["wattleview_mw"]),
        "Bulk-buys counted at what they add by the design point, not systems installed. The seven past campaigns "
        "installed %s systems per 100 houses without PV, but more than half of those households would have gone "
        "solar within three years anyway; read against suburbs whose campaign had not yet run, they added %s per "
        "100 (Solar Streets tab)." % (st["installed_per_100_houses"], st["added_per_100_houses"]),
        "Installed-systems basis puts YRB24 first. On the added basis Bindarra Waters (%s houses without PV) adds "
        "%s MW and YRB24 stays %s MW short of YRB27." % (st["bindarra_houses_without_pv"],
                                                       figs["call"]["bindarra_waters_mw"], figs["call"]["gap_mw"]),
    ]
    for i, t in enumerate(notes, 13):
        ws.cell(row=i, column=1, value=t).font = Font(size=9)
    ws.sheet_view.showGridLines = False

    f = wb.create_sheet("Forecast")
    f["A1"] = "Forecast rooftop PV (micro EG) at 31 August 2029 by feeder"
    f["A1"].font = Font(bold=True, size=11)
    _head(f, 3, ["Feeder", "Connected 31 Aug 2026 (kW)", "Trend (kW a year, FY2023-24 to FY2025-26)",
                 "Base forecast 31 Aug 2029 (kW)", "Systems the 2026 bulk-buy adds", "Bulk-buy addition (kW)",
                 "Forecast 31 Aug 2029 (kW)", "Hosting capacity, 2029 case (MW)",
                 "Above hosting capacity (MW)", "Rank"], [9, 13, 15, 14, 12, 12, 14, 13, 13, 7])
    camp_of = {p["feeder"]: (name, fc["lift"][name]) for name, p in pools.items()}
    for i, fd in enumerate(K.FEEDERS, 4):
        nm = camp_of.get(fd)
        vals = [fd, int(round(c["base"]["cap"].get(fd, 0))), int(round(c["base"]["trend"][fd])),
                int(round(c["base"]["base"][fd])), round(nm[1], 1) if nm else None,
                int(round(fc["add_kw"].get(fd, 0))) if nm else None, int(round(fc["fc"][fd])),
                L["hca"][fd]["hc29"], r1(fc["excess"][fd]), fc["order"].index(fd) + 1]
        for j, v in enumerate(vals, 1):
            cc = f.cell(row=i, column=j, value=v)
            if j in (2, 3, 4, 6, 7):
                cc.number_format = "#,##0"
            elif j == 5:
                cc.number_format = "0.0"
            elif j == 8:
                cc.number_format = "0.00"
            elif j == 9:
                cc.number_format = "0.0"
    tot = [("Yarrabin", None), (sum(c["base"]["cap"].get(fd, 0) for fd in K.FEEDERS), "#,##0"),
           (sum(c["base"]["trend"][fd] for fd in K.FEEDERS), "#,##0"),
           (sum(c["base"]["base"][fd] for fd in K.FEEDERS), "#,##0"),
           (round(sum(fc["lift"].values()), 1), "0.0"), (sum(fc["add_kw"].values()), "#,##0"),
           (sum(fc["fc"][fd] for fd in K.FEEDERS), "#,##0")]
    for j, (v, fmt) in enumerate(tot, 1):
        cc = f.cell(row=12, column=j, value=int(round(v)) if (isinstance(v, float) and fmt == "#,##0") else v)
        cc.font = Font(bold=True)
        if fmt:
            cc.number_format = fmt
    f["A14"] = ("Connected capacity: each NMI's latest connected configuration on the network model's feeder. Trend: "
                "yearly increase in connected capacity less program systems first connected that year (NPP-114 3.2). "
                "Bulk-buy addition: systems the campaign adds by 31 Aug 2029 at the feeder's FY2025-26 average new "
                "inverter (4.1, 4.3). Carried unrounded.")
    f["A14"].font = Font(size=9, italic=True)
    f.freeze_panes = "B4"
    f.page_setup.orientation = "landscape"
    f.page_setup.fitToWidth = 1
    f.sheet_properties.pageSetUpPr.fitToPage = True
    f.print_title_rows = "3:3"

    p = wb.create_sheet("Solar Streets")
    p["A1"] = "Solar Streets: systems installed against systems added, seven past campaigns"
    p["A1"].font = Font(bold=True, size=11)
    _head(p, 3, ["Campaign", "Suburb", "Feeder", "Sign-ups opened", "Window to",
                 "Owner-occupied houses without PV at opening", "Systems installed (program ref)",
                 "First applications, 12 months before", "First applications in window (all, incl. program)",
                 "Comparison suburbs: window / 12 months before", "Expected in window without campaign",
                 "Added by campaign", "Added per 100 houses without PV"],
          [10, 17, 8, 12, 12, 13, 12, 12, 14, 14, 13, 11, 12], height=72)
    for i, x in enumerate(c["camps"], 4):
        rt = c["ratios"][x["key"]]
        vals = [x["key"], x["suburb"], x["feeder"], x["opening"], rt["end"], x["pool"], x["systems"], x["pre"], x["post"],
                round(rt["post"] / rt["pre"], 3), round(x["cf"], 1), round(x["net"], 1),
                round(100.0 * x["net"] / x["pool"], 2)]
        for j, v in enumerate(vals, 1):
            cc = p.cell(row=i, column=j, value=v)
            if j in (4, 5):
                cc.number_format = "d mmm yyyy"
            elif j in (6, 7, 8, 9):
                cc.number_format = "#,##0"
            elif j == 10:
                cc.number_format = "0.000"
            elif j in (11, 12):
                cc.number_format = "#,##0.0"
            elif j == 13:
                cc.number_format = "0.00"
    n = 4 + len(c["camps"])
    p.cell(row=n, column=1, value="All seven").font = Font(bold=True)
    for j, v, fmt in ((6, sum(x["pool"] for x in c["camps"]), "#,##0"), (7, sum(x["systems"] for x in c["camps"]), "#,##0"),
                      (12, round(sum(x["net"] for x in c["camps"]), 1), "#,##0.0"),
                      (13, round(100.0 * c["rate"]["slope"], 2), "0.00")):
        cc = p.cell(row=n, column=j, value=v)
        cc.number_format = fmt
        cc.font = Font(bold=True)
    p.cell(row=n + 1, column=1, value="Installed per 100 houses without PV, all seven: %s" % st["installed_per_100_houses"]).font = \
        Font(size=9)
    m = n + 3
    _head(p, m, ["2026 campaign", "Suburb", "Feeder", "Sign-ups open", "", "Owner-occupied houses without PV, 1 Sep 2026",
                 "Systems added by 31 Aug 2029", "Average new inverter (kW, FY2025-26)", "PV added (MW)"], height=58)
    for i, cm in enumerate(K.CAMPAIGNS_2026, m + 1):
        name = cm["suburb"]
        pl = pools[name]
        vals = [cm["key"], name, pl["feeder"], cm["open"], None, pl["pool"], round(fc["lift"][name], 1),
                round(fc["avgkw"][pl["feeder"]], 2), r1(fc["add_mw"][name])]
        for j, v in enumerate(vals, 1):
            cc = p.cell(row=i, column=j, value=v)
            if j == 4:
                cc.number_format = "d mmm yyyy"
            elif j == 6:
                cc.number_format = "#,##0"
            elif j == 7:
                cc.number_format = "0.0"
            elif j == 8:
                cc.number_format = "0.00"
            elif j == 9:
                cc.number_format = "0.0"
    p.cell(row=m + 4, column=1, value=(
        "Owner-occupied house: rates roll land use 02, rating R1, matched to NMIs via the GIS cross-reference. First "
        "application: earliest application at the house other than a withdrawn one, by received date. Window: 36 months "
        "from sign-ups opening, or to 31 Aug 2026 for BNR23 and PBG24. Comparison: suburbs whose own campaign opens "
        "after the window (not yet treated). Expected without campaign = suburb's first applications in the 12 months "
        "before x comparison suburbs' window / 12-month ratio.")).font = Font(size=9, italic=True)

    d = wb.create_sheet("DER Register")
    d["A1"] = "Active DER Register sites by feeder, extract as at 31 August 2026"
    d["A1"].font = Font(bold=True, size=11)
    _head(d, 3, ["Feeder", "Active sites", "Sites first lodged in FY2025-26", "Inverter capacity (kW)",
                 "Battery storage (kWh)"], [9, 12, 15, 14, 14])
    for i, fd in enumerate(K.FEEDERS, 4):
        vals = [fd, figs["der"][fd]["active_sites"], figs["der"][fd]["first_lodged_fy2025_26"],
                ten(c["asks"]["inverter"][fd]), ten(c["asks"]["storage"][fd])]
        for j, v in enumerate(vals, 1):
            cc = d.cell(row=i, column=j, value=v)
            if j >= 2:
                cc.number_format = "#,##0"
    d["A13"] = ("A site is the premises' NMI (ten characters; a builder's temporary supply NMI read at the NMI that "
                "replaced it), placed on the feeder the network model shows supplying it at 31 August 2026. First "
                "lodged: the earliest registration of any AC connection at the site, decommissioned ones included.")
    d["A14"] = ("Inverter capacity: two-inverter connections lodged on the installer portal before July 2024 carry the "
                "connection's total on each inverter row and are counted once. Battery storage: Sungrow SBR stacks "
                "carry the stack's total on each module row and are counted once; other modular batteries are "
                "lodged module by module and summed. Rounded to the nearest 10 kW and 10 kWh.")
    for cell in ("A13", "A14"):
        d[cell].font = Font(size=9, italic=True)
    d.freeze_panes = "B4"

    n_ = wb.create_sheet("Notes")
    n_["A1"] = "Notes and sources"
    n_["A1"].font = Font(bold=True, size=11)
    lines = [
        ("Base date", "31 August 2026 (NPP-114 rev 3, clause 3.1); connection register extract of 1 September 2026 and "
                      "the network model at 31 August 2026."),
        ("Design point", "31 August 2029, the start of the spring minimum-demand season (clause 2.4)."),
        ("Rooftop PV", "Micro EG of 30 kVA or less at residential and business NMIs, at the inverter's AC rating; "
                       "systems migrated in 2019 take their inverters from the legacy equipment table."),
        ("Hosting capacity", "Buchanan Power Engineering report BPE-2026-031 (24 July 2026), spring 2029 "
                             "minimum-demand case (clause 5.1)."),
        ("Solar Streets 2026", "Merindah City Council letter of 14 August 2026 (reference SS/2026/02): sign-ups "
                               "1 September to 27 November 2026 in Bindarra Waters and Wattleview."),
        ("Past campaigns", "Council Solar Streets summary 2018 to 2024 for sign-up dates; systems from program "
                           "references in the connection register, which match the council's systems installed. What each "
                           "campaign added is read against suburbs whose campaign had not yet run over the same months."),
        ("DER Register", "Extract as at 31 August 2026, AC connections and devices at NMIs on the Yarrabin feeders; "
                         "the NMI re-allocation log for builders' temporary supply NMIs."),
        ("Rounding", "Carried unrounded and rounded once: MW to one decimal (clause 7.1); inverter capacity to the "
                     "nearest 10 kW and battery storage to the nearest 10 kWh."),
    ]
    for i, (k, v) in enumerate(lines, 3):
        n_.cell(row=i, column=1, value=k).font = Font(bold=True, size=10)
        n_.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
        n_.cell(row=i, column=1).alignment = Alignment(vertical="top")
        n_.row_dimensions[i].height = 30
    n_.column_dimensions["A"].width = 22
    n_.column_dimensions["B"].width = 110
    wb.save(path)
    PP.normalize_ooxml(path, K.PEOPLE["analyst"], K.PEOPLE["analyst"], "2026-10-13T00:12:44Z",
                       "2026-10-16T06:41:09Z", excel=True)


def write_chart(path, c, figs):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D
    from matplotlib.patches import FancyBboxPatch, Rectangle
    fc, L = c["fc"], c["L"]
    w = fc["winner"]
    surface, ink, ink2, muted, grid, base_c = "#fcfcfb", "#0b0b0b", "#52514e", "#898781", "#e1e0d9", "#c3c2b7"
    accent, ctx = "#2a78d6", "#c3c2b7"
    plt.rcParams.update({"font.family": "DejaVu Sans", "svg.hashsalt": "task136"})
    fig, ax = plt.subplots(figsize=(10, 5.8), dpi=150)
    fig.patch.set_facecolor(surface)
    ax.set_facecolor(surface)
    vals = [fc["fc"][f] / 1000.0 for f in K.FEEDERS]
    hcs = [L["hca"][f]["hc29"] for f in K.FEEDERS]
    width = 0.34
    for i, (f, v, h) in enumerate(zip(K.FEEDERS, vals, hcs)):
        col = accent if f == w else ctx
        ax.add_patch(FancyBboxPatch((i - width / 2, -1.0), width, v + 1.0, boxstyle="round,pad=0,rounding_size=0.050",
                                    mutation_aspect=4.0, linewidth=0, facecolor=col, zorder=2, clip_on=True))
        ax.plot([i - width / 2 - 0.08, i + width / 2 + 0.08], [h, h], color=ink, linewidth=2.0, solid_capstyle="butt",
                zorder=3)
    ax.set_xlim(-0.6, len(K.FEEDERS) - 0.4)
    top = max(vals) * 1.16
    ax.set_ylim(0, top)
    ax.set_yticks([y for y in range(0, int(top) + 1, 5)])
    ax.set_yticklabels(["%d" % y for y in range(0, int(top) + 1, 5)])
    ax.set_xticks(range(len(K.FEEDERS)))
    ax.set_xticklabels(K.FEEDERS, fontsize=10, color=ink2)
    ax.tick_params(axis="y", colors=muted, labelsize=9, length=0)
    ax.tick_params(axis="x", length=0, pad=6)
    ax.yaxis.grid(True, color=grid, linewidth=0.8, zorder=0)
    ax.set_axisbelow(True)
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.spines["bottom"].set_color(base_c)
    ax.set_ylabel("MW", color=ink2, fontsize=10)
    for f in (w, fc["runner"]):
        i = K.FEEDERS.index(f)
        v = fc["fc"][f] / 1000.0
        ax.annotate(("%s MW above\nhosting capacity\nof %s MW" % (figs["call"]["excess_mw"], figs["steps"]["winner_hosting_mw"]))
                    if f == w else "%s MW above\nhosting capacity" % figs["call"]["runner_up_excess_mw"],
                    xy=(i, v), xytext=(0, 8), textcoords="offset points", ha="center", va="bottom", fontsize=9,
                    color=ink if f == w else ink2)
    fig.text(0.075, 0.955, "%s sits furthest above its hosting capacity at the 31 August 2029 design point" % w,
             fontsize=13, color=ink, weight="bold", ha="left")
    fig.text(0.075, 0.915, "Forecast rooftop PV connected at 31 August 2029 (bars) and each feeder's 2029 hosting "
             "capacity (black marks), Yarrabin 22 kV feeders", fontsize=9.5, color=ink2, ha="left")
    handles = [Rectangle((0, 0), 1, 1, facecolor=accent), Rectangle((0, 0), 1, 1, facecolor=ctx),
               Line2D([0], [0], color=ink, linewidth=2)]
    ax.legend(handles, ["%s forecast PV" % w, "Other feeders' forecast PV", "Hosting capacity, 2029 case"],
              loc="upper right", frameon=False, fontsize=9, labelcolor=ink2, ncol=3, bbox_to_anchor=(1.0, 1.06))
    fig.text(0.075, 0.025, "Source: Mirrabee Energy connection register and network model at 31 August 2026; "
             "hosting capacity assessment BPE-2026-031; NPP-114 rev 3.", fontsize=7.5, color=muted, ha="left")
    fig.subplots_adjust(left=0.075, right=0.985, top=0.84, bottom=0.11)
    fig.savefig(path, facecolor=surface, metadata={"Software": None})
    plt.close(fig)


def main(target, golden, figures_path):
    os.makedirs(golden, exist_ok=True)
    c = compute(target)
    figs = figures(c)
    write_book(os.path.join(golden, K.GOLD_BOOK), c, figs)
    write_chart(os.path.join(golden, K.GOLD_CHART), c, figs)
    with open(figures_path, "w", encoding="utf-8") as fh:
        json.dump(figs, fh, indent=1, ensure_ascii=False, default=str)
        fh.write("\n")
    return c, figs
