"""Independent verifier for task462: reads only the shipped folder and the goldens, on its own code path (pandas,
statsmodels, a k-d tree), and recomputes the rating, each producer's methane emitted and intensity, the gas delivered
to the plant and the gas flared, and the conformed plume file's counts, against generator/figures.json and the golden
workbook. Imports nothing from the generator.

    python3 task462/generator/verify.py
"""
import json
import math
import os
import sys

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
LIMIT = 0.20
KG_PER_MCF = 19.2
RADIUS = 150.0


def find(target, stem):
    hit = [f for f in os.listdir(target) if f.startswith(stem)]
    assert len(hit) == 1, (stem, hit)
    return os.path.join(target, hit[0])


def logit_fit(x, y):
    import statsmodels.api as sm
    res = sm.GLM(y, sm.add_constant(x), family=sm.families.Binomial()).fit()
    return res.params


def recompute(target):
    from scipy.spatial import cKDTree
    reg = pd.read_csv(find(target, "tarvane_tollaver_site_register"), dtype=str)
    reg["lat"] = reg["latitude"].astype(float)
    reg["lon"] = reg["longitude"].astype(float)
    lat0 = math.radians(36.6)
    xy = lambda la, lo: np.column_stack([la * 111_000.0, lo * 111_000.0 * math.cos(lat0)])
    tree = cKDTree(xy(reg["lat"].to_numpy(), reg["lon"].to_numpy()))
    pl = pd.read_csv(find(target, "talmora_plume_detections"), dtype={"nearest_facility": str})
    dist, idx = tree.query(xy(pl["plume_lat"].to_numpy(), pl["plume_lon"].to_numpy()))
    pl["site"] = np.where(dist <= RADIUS, reg["site_id"].to_numpy()[idx], None)
    pl["day"] = pd.to_datetime(pl["image_time_utc"]).dt.date

    # operators over the year: (from, to) tenures per site
    start, end = pd.Timestamp("2025-10-01").date(), pd.Timestamp("2026-10-01").date()
    ten = {}
    for r in reg.itertuples():
        if isinstance(r.previous_operator, str) and r.previous_operator:
            since = pd.Timestamp(r.operator_since).date()
            ten[r.site_id] = [(r.previous_operator, start, since), (r.operator, since, end)]
        else:
            ten[r.site_id] = [(r.operator, start, end)]

    def op_on(site, day):
        for p, a, b in ten[site]:
            if a <= day < b:
                return p
        return ten[site][-1][0]
    pl["producer"] = [op_on(s, d) if isinstance(s, str) else None for s, d in zip(pl["site"], pl["day"])]

    # the detection curve
    rel = pd.read_excel(find(target, "talmora_controlled_release_study"), sheet_name="Releases")
    b0, b1 = logit_fit(np.log(rel["metered_rate_kg_h"].to_numpy()), (rel["plume_detected"] == "Y").astype(float).to_numpy())
    pod = lambda q: 1.0 / (1.0 + np.exp(-(b0 + b1 * np.log(q))))

    # each site's passes from the coverage file through its facility numbers
    cov = pd.read_csv(find(target, "talmora_flight_line_coverage"), dtype=str)
    fac = reg[["site_id", "facility_numbers"]].assign(facility_number=reg["facility_numbers"].str.split("; ")).explode(
        "facility_number")[["site_id", "facility_number"]]
    ps = cov.merge(fac, on="facility_number")[["campaign", "flight_line", "site_id", "imaged"]].drop_duplicates()
    imaged = ps[ps["imaged"] == "Y"]

    camp_day = {"TB-2511": pd.Timestamp("2025-11-10").date(), "TB-2602": pd.Timestamp("2026-02-09").date(),
                "TB-2605": pd.Timestamp("2026-05-11").date(), "TB-2608": pd.Timestamp("2026-08-10").date()}

    def emitted(weighted=True):
        p = pl[pl["site"].notna() & (pl["emission_rate_kg_h"] >= 5.0)].copy()
        p["v"] = p["emission_rate_kg_h"] / (pod(p["emission_rate_kg_h"].to_numpy()) if weighted else 1.0)
        per_pass = p.groupby(["site", "campaign", "flight_line"])["v"].sum()
        im = imaged.copy()
        im["v"] = [per_pass.get((s, c, l), 0.0) for s, c, l in zip(im["site_id"], im["campaign"], im["flight_line"])]
        sc = im.groupby(["site_id", "campaign"])["v"].mean()
        out = {}
        for site, tl in ten.items():
            for prod_, a, b in tl:
                vals = [sc[(site, c)] for c, d0 in camp_day.items() if a <= d0 < b and (site, c) in sc.index]
                if vals:
                    out[prod_] = out.get(prod_, 0.0) + float(np.mean(vals)) * (b - a).days * 24 / 1000.0
        return out

    # methane produced
    ops = pd.read_csv(find(target, "ogc_operator_list"), dtype=str)
    names = sorted(reg["operator"].unique())
    short = {r.operator_number: next((n for n in names if r.operator_name.startswith(n)), None) for r in ops.itertuples()}
    pr = pd.read_csv(find(target, "ogc_well_production"), dtype={"api_number": str})
    pr = pr.sort_values("received_date").groupby(["api_number", "production_month"]).tail(1)
    gas = pd.read_csv(find(target, "tollaver_gas_composite_analyses"))
    gas["year"] = gas["composite_date"].str[:4].astype(int)
    share = gas.sort_values("year").groupby("field_name").tail(1).set_index("field_name")["methane_mol_pct"] / 100.0
    pr["producer"] = pr["operator_number"].map(short)
    pr = pr[pr["producer"].notna()]
    pr["t"] = pr["gas_mcf"] * 14.65 / 14.73 * pr["field_name"].map(share) * KG_PER_MCF / 1000.0
    produced = pr.groupby("producer")["t"].sum().to_dict()

    em = emitted(True)
    em0 = emitted(False)
    inten = {p: 100.0 * em[p] / produced[p] for p in names}
    inten0 = {p: 100.0 * em0[p] / produced[p] for p in names}

    # gas delivered to the plant
    st = pd.read_csv(find(target, "tollaver_plant_allocation_statements"), dtype=str)
    st["rev"] = st["revision"].astype(int)
    st = st[st["rev"] == st.groupby("statement_month")["rev"].transform("max")]
    st["pm"] = st["reference"].str.split().str[1]
    st = st[(st["pm"] >= "2025-10") & (st["pm"] <= "2026-09")]
    hv = pd.read_csv(find(target, "tollaver_plant_station_heating_values"), dtype={"heating_value_btu_scf": float})
    hvm = {(r.station_id, r.month): r.heating_value_btu_scf for r in hv.itertuples()}
    vol = st["volume_mcf"].astype(float)
    conv = [float(e) * 1000.0 / hvm[(s, m)] for s, m, e in zip(st["station_id"], st["pm"], st["energy_mmbtu"])]
    st["mcf"] = np.where(vol.notna(), vol, conv)
    stn = pd.read_csv(find(target, "tollaver_plant_meter_stations"), dtype=str)
    st["producer"] = st["station_id"].map(dict(zip(stn["station_id"], stn["operator_number"]))).map(short)
    delivered = (st[st["producer"].notna()].groupby("producer")["mcf"].sum() / 1000.0).to_dict()

    # gas flared
    fl = pd.read_csv(find(target, "ogc_flaring_venting_reports"), dtype=str)
    fl["mcf"] = fl["volume_mcf"].astype(float)
    fl = fl.sort_values("received_date").groupby(["lease_number", "production_month", "disposition_code"]).tail(1)
    fl = fl[fl["disposition_code"].isin(["FL", "FR"]) & (fl["production_month"] >= "2025-10") &
            (fl["production_month"] <= "2026-09")]
    lr = pd.read_csv(find(target, "ogc_lease_register"), dtype=str).sort_values("effective_date")

    def lease_op(lease, month):
        h = lr[lr["lease_number"] == lease]
        h = h[h["effective_date"] <= month + "-01"]
        return short.get(h["operator_number"].iloc[-1])
    lease_ops = {}
    fl["producer"] = [lease_ops.setdefault((l, m), lease_op(l, m)) for l, m in zip(fl["lease_number"], fl["production_month"])]
    flared = (fl[fl["producer"].notna()].groupby("producer")["mcf"].sum() / 1000.0).to_dict()
    return dict(names=names, inten=inten, inten0=inten0, em=em, delivered=delivered, flared=flared,
                rated=sorted(p for p in names if inten[p] <= LIMIT), rated0=sorted(p for p in names if inten0[p] <= LIMIT),
                plumes=len(pl), attributed=int(pl["site"].notna().sum()), q50=math.exp(-b0 / b1),
                q90=math.exp((math.log(9) - b0) / b1))


def main(target, golden, fig_path=None):
    F = json.load(open(fig_path or os.path.join(HERE, "figures.json")))
    V = recompute(target)
    bad = []

    def check(label, cond, detail=""):
        print("%s  verify: %s%s" % ("PASS" if cond else "FAIL", label, ("  [" + detail + "]") if detail else ""))
        if not cond:
            bad.append(label)
    check("the rated producers", V["rated"] == sorted(F["rated"]), ", ".join(V["rated"]))
    check("each producer's methane emitted, whole tonnes", all(round(V["em"][p]) == F["emitted_t"][p] for p in V["names"]))
    check("each producer's measured intensity, three decimals", all(round(V["inten"][p], 3) == F["intensity"][p]
                                                                    for p in V["names"]))
    check("the default path (blanks as zeros) rates eight", len(V["rated0"]) == 8, ", ".join(V["rated0"]))
    check("the detection curve's 50 and 90 per cent rates", round(V["q50"], 1) == F["q50"] and round(V["q90"], 1) == F["q90"],
          "%.2f, %.2f" % (V["q50"], V["q90"]))
    check("gas delivered to the plant, whole MMcf", all(round(V["delivered"][p]) == F["delivered_mmcf"][p] for p in V["names"]))
    check("gas flared, whole MMcf", all(round(V["flared"][p]) == F["flared_mmcf"][p] for p in V["names"]))
    check("plumes and plumes with a producer site", V["plumes"] == F["plumes"] and V["attributed"] == F["plumes_attributed"],
          "%d, %d" % (V["plumes"], V["attributed"]))
    import openpyxl
    ws = openpyxl.load_workbook(os.path.join(golden, "tollaver_producer_methane_ratings_2027.xlsx"))["2027 ratings"]
    rows = {ws.cell(row=r, column=1).value: [ws.cell(row=r, column=c).value for c in range(3, 8)] for r in range(5, 17)}
    check("the golden workbook's rows", all(rows[p][0] == round(V["em"][p]) and abs(rows[p][1] - round(V["inten"][p], 3)) < 1e-9
                                             and rows[p][3] == round(V["delivered"][p]) and rows[p][4] == round(V["flared"][p])
                                             for p in V["names"]))
    conf = pd.read_csv(os.path.join(golden, "talmora_plumes_conformed_2025_2026.csv"), dtype=str)
    check("the conformed plume file", len(conf) == V["plumes"] and int(conf["site_id"].notna().sum()) == V["attributed"])
    if bad:
        raise SystemExit("verify failed: %s" % bad)
    return V


if __name__ == "__main__":
    main(os.path.join(TASK, "target"), os.path.join(TASK, "golden"))
