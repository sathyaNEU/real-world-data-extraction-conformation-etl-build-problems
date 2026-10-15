"""The golden deliverables, computed from the shipped pack alone: the 2027 ratings workbook (each producer's methane
emitted, measured intensity and rating, the gas it delivered to the plant and flared, the intensity chart, the detection
curve and the notes) and the conformed plume file; and generator/figures.json, every figure the write-up states."""
import csv
import io
import json
import math
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import bundle as BD    # noqa: E402
import methods as M    # noqa: E402
import world as W      # noqa: E402

ORDER = sorted(W.NAMES)


def compute(target):
    B = BD.load(target)
    sites, passes, pl = B["sites"], B["passes"], B["plumes"]
    site = {s["site_id"]: s for s in sites}
    for p, (sid, d) in zip(pl, M.attribute(pl, sites)):
        p["site"], p["dist"] = sid, d
        p["producer"] = M.operator_on(site[sid], p["time"].date()) if sid else None
    pod = M.fit_pod(B["releases"])
    w = lambda z: z["rate"] / pod(z["rate"])
    prod = M.produced(B["production"], B["composites"])
    em = M.emitted(sites, passes, pl, w)
    inten = M.intensity(em, prod)
    R = {"R0 reported 2025 inventories": M.inventory(B["ghgrp"]),
         "R1 each detected source carried through the year": M.carried_sources(sites, passes, pl),
         "R2 the campaigns averaged, blanks as zeros": M.emitted(sites, passes, pl, lambda z: z["rate"]),
         "R3 detection-weighted, the campaigns averaged": em,
         "x each source at the share of campaigns it was detected in": M.carried_sources(sites, passes, pl, duty=True)}
    RI = {k: M.intensity(v, prod) for k, v in R.items()}
    st, sn, hv = B["statements"], B["stations"], B["heating"]
    deliv = M.delivered(st, sn, hv)
    flar = M.flared(B["flaring"], B["leases"])
    img = {}
    for ps in passes:
        img.setdefault((ps["site"], ps["campaign"]), []).append(ps["imaged"])
    return dict(B=B, pod=pod, prod=prod, em=em, inten=inten, rated=M.rated(inten), R=R, RI=RI, deliv=deliv, flar=flar,
                plumes=pl, site=site, img=img)


def curve(pod):
    b0, b1 = float(pod.b[0]), float(pod.b[1])
    return dict(b0=b0, b1=b1, q50=math.exp(-b0 / b1), q90=math.exp((math.log(9.0) - b0) / b1))


def numbers(target):
    """Operator number for each rated producer, from the commission's list."""
    return {v: k for k, v in BD.operators(target).items() if v in W.NAMES}


def write_csv(X, path):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(["plume_id", "campaign", "flight_line", "image_time_utc", "plume_lat", "plume_lon", "emission_rate_kg_h",
                "site_id", "site_type", "producer", "distance_to_site_m"])
    for p in X["plumes"]:
        r = p["raw"]
        s = X["site"].get(p["site"])
        w.writerow([r["plume_id"], r["campaign"], r["flight_line"], r["image_time_utc"], r["plume_lat"], r["plume_lon"],
                    r["emission_rate_kg_h"], p["site"] or "", s["kind"] if s else "", p["producer"] or "",
                    "" if p["dist"] is None else "%d" % round(p["dist"])])
    with open(path, "w", newline="", encoding="utf-8") as f:
        f.write(buf.getvalue())
    return len(X["plumes"])


def write_workbook(X, path, target):
    import openpyxl
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    em, inten, deliv, flar = X["em"], X["inten"], X["deliv"], X["flar"]
    rated = sorted(X["rated"], key=lambda p: inten[p])
    out = sorted((p for p in W.NAMES if p not in X["rated"]), key=lambda p: inten[p])
    cv = curve(X["pod"])
    num = numbers(target)
    wb = openpyxl.Workbook()
    bold = Font(name="Calibri", size=11, bold=True)
    title = Font(name="Calibri", size=14, bold=True, color="1F4E3D")
    grey = Font(name="Calibri", size=9, italic=True, color="5A5F66")
    head = PatternFill("solid", fgColor="DDEBE3")
    pick = PatternFill("solid", fgColor="EAF4EC")
    thin = Border(bottom=Side(style="thin", color="9AA0A6"))
    top = Border(top=Side(style="thin", color="5A5F66"))

    ws = wb.active
    ws.title = "2027 ratings"
    ws["A1"] = "Tollaver Basin 2027 low-methane ratings: %s" % ", ".join(rated[:-1]) + " and " + rated[-1]
    ws["A1"].font = title
    ws["A2"] = ("Measured methane intensity over the rating year 1 October 2025 to 30 September 2026, TMR-STD-03 v1.1: "
                "rated low-methane at 0.20 per cent or less (3.1). Methane in metric tonnes; gas in million cubic feet "
                "at 14.73 psia.")
    ws["A2"].font = grey
    ws.merge_cells("A2:G2")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 30
    hdr = ["Producer", "Operator no.", "Methane emitted (t)", "Measured methane intensity (%)", "2027 rating",
           "Gas delivered to the plant (MMcf)", "Gas flared (MMcf)"]
    ws.append([])
    ws.append(hdr)
    for j in range(1, 8):
        c = ws.cell(row=4, column=j)
        c.font = bold
        c.fill = head
        c.border = thin
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 32
    for p in ORDER:
        ok = p in X["rated"]
        ws.append([p, num[p], round(em[p]), round(inten[p], 3), "Low-methane" if ok else "Not rated", round(deliv[p]),
                   round(flar[p])])
        r = ws.max_row
        ws.cell(row=r, column=3).number_format = "#,##0"
        ws.cell(row=r, column=4).number_format = "0.000"
        ws.cell(row=r, column=6).number_format = "#,##0"
        ws.cell(row=r, column=7).number_format = "#,##0"
        ws.cell(row=r, column=2).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=5).alignment = Alignment(horizontal="center")
        if ok:
            for j in range(1, 8):
                ws.cell(row=r, column=j).fill = pick
            ws.cell(row=r, column=5).font = bold
    ws.append(["Twelve producers", None, round(sum(em.values())), None, "%d rated" % len(X["rated"]),
               round(sum(deliv.values())), round(sum(flar.values()))])
    r = ws.max_row
    for j in range(1, 8):
        ws.cell(row=r, column=j).font = bold
        ws.cell(row=r, column=j).border = top
    for j in (3, 6, 7):
        ws.cell(row=r, column=j).number_format = "#,##0"
    ws.cell(row=r, column=5).alignment = Alignment(horizontal="center")
    ws.append([])
    near_in, near_out = rated[-1], out[0]
    B = X["B"]
    rel = B["releases"]
    notes = [
        "Rated: %s. Nearest the line: %s at %.3f per cent among the rated, %s at %.3f per cent among the rest." % (
            "; ".join("%s %.3f" % (p, inten[p]) for p in rated), near_in, inten[near_in], near_out, inten[near_out]),
        "Methane emitted (TMR-STD-03 2.5, 4.2, 4.3): each plume Talmora reported at 5 kg/h or more, attributed to the "
        "register site nearest it within 150 m, counted at its rate divided by the chance that a pass detects a plume "
        "of that rate, since a pass does not detect every plume (Talmora TA-TN-26-07, 2.1). The chance is Talmora's "
        "logistic curve in the log of the rate fitted by maximum likelihood to the %d controlled releases (%d "
        "detected): 50 per cent at %.1f kg/h, 90 per cent at %.1f kg/h. A site's rate on a pass is the sum of its "
        "counted plumes, in a campaign the mean of the passes that imaged it, and over each operator's part of the "
        "year the mean of its campaigns in that part, times that part's hours." % (len(rel), sum(r["detected"] for r in rel), cv["q50"],
                                                                 cv["q90"]),
        "Methane produced (2.3, 2.4): the latest report for each well and month, to the operator of record that month, "
        "gas at the commission's 14.65 psia restated to 14.73, times the field's 2026 composite methane share and 19.2 "
        "kg per Mcf.",
        "Gas delivered: the plant's statement lines for production months October 2025 to September 2026, the month "
        "taken from each line's reference (prior-period adjustments booked in later statements included), the March "
        "2026 statement at revision 2, and the two stations metering energy only from July 2026 restated to Mcf at "
        "each station's heating value for the month; stations to producers through the plant's station list.",
        "Gas flared: the commission's reports for October 2025 to September 2026 under codes FL and, from January 2026, "
        "FR (routine flaring, which the commission's basin summary counts as flared), the latest report for each lease, "
        "month and code, each lease to its operator of record in the month (lease register).",
    ]
    for t in notes:
        ws.append([t])
        c = ws.cell(row=ws.max_row, column=1)
        c.font = grey
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=7)
        ws.row_dimensions[ws.max_row].height = 16 + 13 * (len(t) // 120)
    for col, wd in zip("ABCDEFG", (27, 12, 13, 15, 13, 17, 12)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "B5"
    ws.print_title_rows = "4:4"

    ws2 = wb.create_sheet("Intensity chart")
    ws2["A1"] = "Measured methane intensity by producer, rating year to 30 September 2026, against the 0.20 per cent line"
    ws2["A1"].font = title
    ws2.append([])
    ws2.append(["Producer", "Measured methane intensity (%)", "Rating line (%)"])
    for j in range(1, 4):
        c = ws2.cell(row=3, column=j)
        c.font = bold
        c.fill = head
        c.border = thin
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    order = rated + out
    for p in order:
        ws2.append([p, round(inten[p], 3), W.LIMIT])
        ws2.cell(row=ws2.max_row, column=2).number_format = "0.000"
        ws2.cell(row=ws2.max_row, column=3).number_format = "0.00"
    n = len(order)
    bar = BarChart()
    bar.type = "col"
    bar.title = "Four producers at or under 0.20 per cent"
    bar.y_axis.title = "Measured methane intensity (%)"
    bar.x_axis.title = "Producer"
    bar.add_data(Reference(ws2, min_col=2, min_row=3, max_row=3 + n), titles_from_data=True)
    cats = Reference(ws2, min_col=1, min_row=4, max_row=3 + n)
    bar.set_categories(cats)
    bar.dataLabels = DataLabelList()
    bar.dataLabels.showVal = True
    bar.y_axis.scaling.min = 0
    bar.y_axis.scaling.max = 0.45
    bar.y_axis.majorUnit = 0.05
    bar.y_axis.numFmt = "0.00"
    bar.legend.position = "b"
    line = LineChart()
    line.add_data(Reference(ws2, min_col=3, min_row=3, max_row=3 + n), titles_from_data=True)
    line.set_categories(cats)
    bar += line
    bar.height = 10
    bar.width = 24
    ws2.add_chart(bar, "E3")
    ws2.column_dimensions["A"].width = 27
    ws2.column_dimensions["B"].width = 17
    ws2.column_dimensions["C"].width = 11

    ws3 = wb.create_sheet("Detection curve")
    ws3["A1"] = "Talmora's chance of detecting a plume, fitted to its 2025 controlled releases"
    ws3["A1"].font = title
    ws3.append([])
    ws3.append(["Logistic curve in ln(rate, kg/h), maximum likelihood", None])
    ws3.cell(row=3, column=1).font = bold
    for label, v, fmt in (("Intercept", round(cv["b0"], 4), "0.0000"), ("Slope on ln(rate)", round(cv["b1"], 4), "0.0000"),
                          ("Rate detected half the time (kg/h)", round(cv["q50"], 1), "0.0"),
                          ("Rate detected nine times in ten (kg/h)", round(cv["q90"], 1), "0.0")):
        ws3.append([label, v])
        ws3.cell(row=ws3.max_row, column=2).number_format = fmt
    ws3.append([])
    ws3.append(["Plume rate (kg/h)", "Chance a pass detects it (%)"])
    for j in (1, 2):
        c = ws3.cell(row=ws3.max_row, column=j)
        c.font = bold
        c.fill = head
        c.border = thin
    for q in (5, 7.5, 10, 15, 20, 30, 50, 100):
        ws3.append([q, round(100 * X["pod"](q), 1)])
        ws3.cell(row=ws3.max_row, column=1).number_format = "0.0"
        ws3.cell(row=ws3.max_row, column=2).number_format = "0.0"
    ws3.append([])
    ws3.append(["Source: %s, sheet Releases (%d releases, %d detected)." % (W.FILES["releases"], len(rel),
                                                                           sum(r["detected"] for r in rel))])
    ws3.cell(row=ws3.max_row, column=1).font = grey
    ws3.column_dimensions["A"].width = 44
    ws3.column_dimensions["B"].width = 26

    ws4 = wb.create_sheet("Sources")
    rows = [["File", "Used for"],
            [W.FILES["standard"], "the rating, the measured methane intensity, sites, wells and the campaign average"],
            [W.FILES["plumes"], "each plume's origin, campaign, flight line and rate"],
            [W.FILES["coverage"], "the sites each flight line imaged, and the lines lost to cloud"],
            [W.FILES["sites"], "each site's location, facility numbers and operators"],
            [W.FILES["releases"], "the detection curve"],
            [W.FILES["methods"], "how Talmora detects and quantifies plumes, and its detection curve's form"],
            [W.FILES["production"], "each well's gas by month, the operator of record and the latest report"],
            [W.FILES["operators"], "operator numbers and names"],
            [W.FILES["gas"], "each field's methane share, 2026 composites"],
            [W.FILES["statements"], "gas delivered to the plant"],
            [W.FILES["stations"], "each station's producer and metering"],
            [W.FILES["heating"], "heating values for the energy-only stations"],
            [W.FILES["flaring"], "gas flared by lease and month"],
            [W.FILES["leases"], "each lease's operator of record by month"],
            [W.FILES["basin"], "the commission's basin totals of flared gas"]]
    for r in rows:
        ws4.append(r)
    for j in (1, 2):
        ws4.cell(row=1, column=j).font = bold
        ws4.cell(row=1, column=j).fill = head
    ws4.column_dimensions["A"].width = 54
    ws4.column_dimensions["B"].width = 78
    ws4.append([])
    ws4.append(["Prepared by %s, Head of Methane Ratings, for the ratings committee of 9 December 2026." %
                W.PEOPLE["ratings"]])
    ws4.cell(row=ws4.max_row, column=1).font = grey
    wb.save(path)


def figures(X, csv_rows, target):
    em, inten, prod, deliv, flar = X["em"], X["inten"], X["prod"], X["deliv"], X["flar"]
    rated = sorted(X["rated"], key=lambda p: inten[p])
    out = sorted((p for p in W.NAMES if p not in X["rated"]), key=lambda p: inten[p])
    cv = curve(X["pod"])
    rel = X["B"]["releases"]
    pl = X["plumes"]
    at = [p for p in pl if p["site"]]
    R2 = X["RI"]["R2 the campaigns averaged, blanks as zeros"]
    img = X["img"]
    return dict(
        rated=rated, unrated=out, limit=W.LIMIT,
        nearest_in=rated[-1], nearest_out=out[0],
        emitted_t={p: int(round(em[p])) for p in W.NAMES},
        emitted_unrounded=em,
        intensity={p: round(inten[p], 3) for p in W.NAMES},
        intensity_unrounded=inten,
        produced_t={p: int(round(prod[p])) for p in W.NAMES},
        delivered_mmcf={p: int(round(deliv[p])) for p in W.NAMES},
        delivered_unrounded=deliv,
        flared_mmcf={p: int(round(flar[p])) for p in W.NAMES},
        flared_unrounded=flar,
        emitted_total_t=int(round(sum(em.values()))),
        delivered_total_mmcf=int(round(sum(deliv.values()))),
        flared_total_mmcf=int(round(sum(flar.values()))),
        q50=round(cv["q50"], 1), q90=round(cv["q90"], 1), b0=cv["b0"], b1=cv["b1"],
        releases=len(rel), releases_detected=sum(r["detected"] for r in rel),
        plumes=len(pl), plumes_attributed=len(at), plumes_at_floor=sum(1 for p in at if p["rate"] >= W.FLOOR),
        plumes_unattributed=len(pl) - len(at),
        site_campaigns_two_passes=sum(1 for v in img.values() if len(v) == 2 and all(v)),
        site_campaigns_unimaged=sum(1 for v in img.values() if not any(v)),
        sites=len(X["site"]),
        rungs={k: dict(rated=M.rated(v), intensity={p: round(v[p], 3) for p in W.NAMES}) for k, v in X["RI"].items()},
        default_wrongly_rated=sorted((p for p in M.rated(R2) if p not in X["rated"]), key=lambda p: R2[p]),
        csv_rows=csv_rows,
        radius_m=int(W.ATTRIB_M), floor=int(W.FLOOR), state_psia=W.STATE_PSIA, std_psia=W.STD_PSIA,
        kg_per_mcf=W.KG_PER_MCF, hours=W.HOURS,
    )


def main(target, out_dir, fig_path):
    os.makedirs(out_dir, exist_ok=True)
    X = compute(target)
    write_workbook(X, os.path.join(out_dir, W.GOLDEN["workbook"]), target)
    n = write_csv(X, os.path.join(out_dir, W.GOLDEN["plumes"]))
    F = figures(X, n, target)
    with open(fig_path, "w") as f:
        json.dump(F, f, indent=1, sort_keys=True)
        f.write("\n")
    return X, F
