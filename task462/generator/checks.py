"""task462 checks: the three rules executed (R1 to R3c), every semantic fork (SF1 to SF13), the ladder, the asks'
devices and stops, the separation line, the distractors, the input gates and the realism sweep.

    python3 task462/generator/checks.py      (run by build.py after the build)
"""
import datetime as dt
import json
import math
import os
import re
import shutil
import sys
import tempfile
import zipfile

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import basin as BA      # noqa: E402
import bundle as BD     # noqa: E402
import golden as GO     # noqa: E402
import methods as M     # noqa: E402
import world as W       # noqa: E402

TASK = os.path.dirname(HERE)
RESULTS = []
FOUR = {"Calbrent Energy", "Brushy Mesa Operating", "Corrandale Petroleum", "Ostrel Energy"}
SMALL_LEAK = {"Pinyon Draw Production", "Anchell Oil & Gas", "Pellisto Gas Partners", "Ashmund Production"}
TANK = {"Brushy Mesa Operating", "Corrandale Petroleum", "Ruskmoor Exploration"}


def ok(label, cond, detail=""):
    RESULTS.append((label, bool(cond), detail))


def edge(v, unit):
    """Distance from the rounding edge at the given unit."""
    x = v / unit
    return abs(x - math.floor(x) - 0.5) * unit


def names(s):
    return " ".join(sorted(p.split()[0] for p in s)) or "none"


def main_call(target, attribution="nearest", passes_mode="register", **kw):
    """The rating from the main call's files alone: sites, coverage, plumes, releases, production, composites."""
    st = BD.sites(target)
    if passes_mode == "one_to_one":
        first = {}
        for s in sorted(st, key=lambda z: z["site_id"]):
            for f in s["fac"]:
                first.setdefault(f, s["site_id"])
        seen = {}
        for r in BD._rows(target, "coverage"):
            sid = first.get(r["facility_number"])
            if sid and (r["campaign"], r["flight_line"], sid) not in seen:
                seen[(r["campaign"], r["flight_line"], sid)] = dict(campaign=r["campaign"], line=r["flight_line"],
                                                                  time=BD._t(r["line_start_utc"]), site=sid,
                                                                  imaged=r["imaged"] == "Y")
        ps = list(seen.values())
    else:
        ps = BD.passes(target, st)
    pl = BD.plumes(target)
    if attribution == "nearest":
        for p, (sid, _) in zip(pl, M.attribute(pl, st)):
            p["site"] = sid
    else:
        for p, sid in zip(pl, M.attribute_by_facility(pl, st)):
            p["site"] = sid
    pod = kw.pop("pod", None) or M.fit_pod(BD.releases(target))
    prod = M.produced(BD.production(target), BD.composites(target), **kw.pop("prod_kw", {}))
    em = M.emitted(st, ps, pl, lambda z: z["rate"] / pod(z["rate"]), **kw)
    return M.intensity(em, prod), em, prod


def asks(target):
    return (M.delivered(BD.statements(target), BD.stations(target), BD.heating(target)),
            M.flared(BD.flaring(target), BD.leases(target)))


def glm_pod(releases, link, x="log"):
    import statsmodels.api as sm
    q = np.array([r["rate"] for r in releases])
    y = np.array([1.0 if r["detected"] else 0.0 for r in releases])
    L = {"cloglog": sm.families.links.CLogLog(), "logit": sm.families.links.Logit()}[link]
    xs = np.log(q) if x == "log" else q
    res = sm.GLM(y, np.column_stack([np.ones(len(q)), xs]), family=sm.families.Binomial(link=L)).fit()
    return lambda rate, wind_ms=None: float(res.predict(np.array([[1.0, math.log(rate) if x == "log" else rate]]))[0])


def run(target, golden, x=None, meta_path=None, fig_path=None):
    F = json.load(open(fig_path or os.path.join(HERE, "figures.json")))
    meta = json.load(open(meta_path or os.path.join(TASK, "metadata.json")))
    files = sorted(os.listdir(target))
    X = GO.compute(target)
    B = X["B"]
    sites, passes, pl = B["sites"], B["passes"], B["plumes"]
    inten, em, prod, pod = X["inten"], X["em"], X["prod"], X["pod"]
    w = lambda z: z["rate"] / pod(z["rate"])
    gold = set(X["rated"])
    RI = X["RI"]
    from pypdf import PdfReader
    doc_text = {f: re.sub(r"\s+", " ", " ".join(p.extract_text() for p in PdfReader(os.path.join(target, f)).pages))
                for f in files if f.endswith(".pdf")}
    std, note = doc_text[W.FILES["standard"]], doc_text[W.FILES["methods"]]

    # ------------------------------------------------------------------ the three rules (dataset-generation 7.1)
    import openpyxl
    heads = {f: open(os.path.join(target, f), encoding="utf-8").readline().lower() for f in files if f.endswith(".csv")}
    rel_ws = openpyxl.load_workbook(os.path.join(target, W.FILES["releases"]))["Releases"]
    heads[W.FILES["releases"]] = ",".join(str(c.value).lower() for c in rel_ws[1])
    ok("R1 the golden reads the raw plume, coverage, release and register rows",
       len(pl) == F["plumes"] and len(B["releases"]) == F["releases"] and len(sites) == F["sites"] and
       len(passes) > 3000, "%d plumes, %d releases, %d sites, %d site passes" % (len(pl), len(B["releases"]),
                                                                              len(sites), len(passes)))
    ok("R1 no shipped column holds a detection probability, a weight, an emitted total, an intensity or a rating",
       not any(re.search(r"probab|weight|pod\b|intensity|emitted|annual|rating", h) for h in heads.values()))
    x = x or __import__("build").world()
    comp_c = {f: {2026: v} for f, v in BA.FIELD_CH4.items()}
    prod_c = M.produced(x["rows"], comp_c)
    st_c, ps_c, pl_c = x["sites"], x["passes"], x["plumes"]
    e2c = M.intensity(M.emitted(st_c, ps_c, pl_c, lambda z: z["rate"]), prod_c)
    ok("R3a the default path on the clean truth (every field read, every filed rule applied) rates the eight",
       set(M.rated(e2c)) == FOUR | SMALL_LEAK, names(M.rated(e2c)))
    pod_c = M.fit_pod(x["releases"])
    e3c = M.intensity(M.emitted(st_c, ps_c, pl_c, lambda z: z["rate"] / pod_c(z["rate"])), prod_c)
    ok("R3b the default path with only the decisive step changed (each plume over its detection chance) rates the four",
       set(M.rated(e3c)) == FOUR, names(M.rated(e3c)))
    e1c = M.intensity(M.carried_sources(st_c, ps_c, pl_c), prod_c)
    e4c = M.intensity(M.carried_sources(st_c, ps_c, pl_c, duty=True), prod_c)
    ok("R3b each lower correction alone stays wrong: sources carried through the year rate none, sources at their "
       "detected share of campaigns rate ten", not M.rated(e1c) and len(M.rated(e4c)) == 10,
       "%s / %s" % (names(M.rated(e1c)), names(M.rated(e4c))))
    tr = M.intensity(x["truth"], prod_c)
    ok("R3 the sites' true time-averaged emissions rate the same four", set(M.rated(tr)) == FOUR, names(M.rated(tr)))
    prod_rows = BD.production(target)
    series = {}
    for r in prod_rows:
        if r["kind"] == "ORIGINAL":
            series.setdefault(r["api"], {})[(r["y"], r["m"])] = r["gas"]
    vecs = [tuple(v[k] for k in sorted(v)) for v in series.values() if len(v) >= 6]
    ok("R3c no well's monthly gas repeats another's", len(set(vecs)) == len(vecs), "%d wells" % len(vecs))
    trip = [(p["rate"], p["lat"], p["lon"]) for p in pl]
    ok("R3c no plume repeats another's rate and origin", len(set(trip)) == len(trip))
    fl_series = {}
    for r in BD.flaring(target):
        if r["kind"] == "ORIGINAL" and r["code"] == "VT":
            fl_series.setdefault(r["lease"], []).append(r["mcf"])
    fv = [tuple(v) for v in fl_series.values() if len(v) >= 6]
    ok("R3c no lease's vented series repeats another's", len(set(fv)) == len(fv))
    alltext = " ".join(doc_text.values())
    ok("R3c no document names a producer, a rating, an intensity figure, a weight or a correction",
       not any(p.split()[0] in alltext for p in W.NAMES) and
       not re.search(r"weight|divid|correct|inverse|Horvitz|per cent at \d", alltext, re.I))
    ok("R3c the standard files the estimand: emitted at 5 kg/h or more over the rating year against methane produced, "
       "0.20 per cent or less", "5 kg per hour or more" in std and "0.20 per cent or less" in std and
       "as a share of its methane produced" in std and "1 October to 30 September" in std)
    ok("R3c the provider's note says a pass does not detect every plume and gives its detection curve's form, no level",
       "does not detect every plume" in note and "logistic function of the natural logarithm" in note and
       not re.search(r"\d+(\.\d+)? ?kg/h (is )?detected|50 per cent|90 per cent", note))

    # ------------------------------------------------------------------------------------------- the ladder
    rated = sorted(gold, key=lambda p: inten[p])
    unrated = sorted((p for p in W.NAMES if p not in gold), key=lambda p: inten[p])
    ok("golden: Calbrent, Brushy Mesa, Corrandale and Ostrel", gold == FOUR, names(gold))
    ok("golden: the rated producer nearest the line 25 per cent or more under it, the nearest unrated 10 per cent or "
       "more over", inten[rated[-1]] <= 0.75 * W.LIMIT and inten[unrated[0]] >= 1.10 * W.LIMIT,
       "%s %.3f, %s %.3f" % (rated[-1], inten[rated[-1]], unrated[0], inten[unrated[0]]))
    r2 = RI["R2 the campaigns averaged, blanks as zeros"]
    ok("R2 (the default path) rates eight: the four and Pinyon Draw, Anchell, Pellisto and Ashmund",
       set(M.rated(r2)) == FOUR | SMALL_LEAK, names(M.rated(r2)))
    ok("R2 the four it rates wrongly sit 9 per cent or more under the line", all(r2[p] <= 0.91 * W.LIMIT for p in SMALL_LEAK),
       ", ".join("%s %.3f" % (p.split()[0], r2[p]) for p in sorted(SMALL_LEAK)))
    r1 = RI["R1 each detected source carried through the year"]
    ok("R1 (each detected source carried through the year) rates none, the lowest 5 per cent or more over",
       not M.rated(r1) and min(r1.values()) >= 1.05 * W.LIMIT, "lowest %.3f" % min(r1.values()))
    rx = RI["x each source at the share of campaigns it was detected in"]
    ok("x (each source at its detected share of campaigns) rates ten", len(M.rated(rx)) == 10, names(M.rated(rx)))
    r0 = RI["R0 reported 2025 inventories"]
    ok("R0 (the producers' reported 2025 inventories) rates all twelve", len(M.rated(r0)) == 12,
       "highest %.3f" % max(r0.values()))
    r1w = M.intensity(M.carried_sources(sites, passes, pl, weight=w), prod)
    ok("grid: sources carried with the weights rate none", not M.rated(r1w))
    agg = {}
    for p_ in W.NAMES:
        mine = [q for q in pl if q["site"] and q["rate"] >= W.FLOOR and
                M.operator_on(X["site"][q["site"]], q["time"].date()) == p_]
        agg[p_] = r2[p_] / float(np.mean([pod(q["rate"]) for q in mine]))
    ok("convergence: the correction applied coarsely (each producer's default total over the mean detection chance of "
       "its plumes) rates the same four", set(M.rated(agg)) == gold, names(M.rated(agg)))
    leak_share = {}
    owner = {z["site_id"]: max(M.tenures(z), key=lambda t: (t[2] - t[1]).days)[0] for z in x["sites"]}
    for p_ in W.NAMES:
        mine = [k for k in x["sources"] if k["rate"] >= W.FLOOR and owner[k["site"]] == p_]
        tot = sum(k["rate"] * k["duty"] for k in mine if k["kind"] != "episode") + 1e-9
        leak_share[p_] = sum(k["rate"] for k in mine if k["kind"] == "leak") / tot
    ok("dominance: Pinyon Draw, Anchell and Pellisto hold 85 per cent or more of their methane outside episodes in "
       "persistent leaks, Ashmund 60 per cent or more, the tank-dump producers 20 per cent or less",
       all(leak_share[p] >= 0.85 for p in SMALL_LEAK - {"Ashmund Production"}) and
       leak_share["Ashmund Production"] >= 0.60 and all(leak_share[p] <= 0.20 for p in TANK), ", ".join("%s %.2f" % (p.split()[0], leak_share[p]) for p in
                                                           sorted(SMALL_LEAK | TANK)))
    up = {p: inten[p] / r2[p] - 1 for p in W.NAMES}
    ok("worth: the weighting moves the leak producers up 40 per cent or more and the tank-dump producers 12 per cent or "
       "less", all(up[p] >= 0.40 for p in SMALL_LEAK - {"Ashmund Production"}) and all(up[p] <= 0.12 for p in TANK),
       ", ".join("%s %+.0f%%" % (p.split()[0], 100 * up[p]) for p in sorted(SMALL_LEAK | TANK)))
    sat = B["satellite"]["estimates"]
    lo = sum(e["ci95_low"] for e in sat) / len(sat) * W.HOURS
    hi = sum(e["ci95_high"] for e in sat) / len(sat) * W.HOURS
    tots = {k: sum(X["R"][k].values()) for k in X["R"] if not k.startswith("R0")}
    ok("oracle: the satellite's basin range holds every rung's total for the twelve", all(lo <= v <= hi for v in tots.values()),
       "%.0f to %.0f t; rungs %s" % (lo, hi, ", ".join("%.0f" % v for v in tots.values())))
    ok("data-native: the releases detect 2 or fewer of those under 3 kg/h and all but one or none over 40 kg/h",
       sum(r["detected"] for r in B["releases"] if r["rate"] < 3) <= 2 and
       sum(not r["detected"] for r in B["releases"] if r["rate"] > 40) <= 1)

    # ---------------------------------------------------------------------------------------- the forks
    def same(label, iv, detail=""):
        ok(label, set(M.rated(iv)) == gold, detail or names(M.rated(iv)))
    i_cur = M.intensity(M.emitted(sites, passes, pl, w, by_current_operator=True),
                        M.produced(B["production"], B["composites"], current_operator=True))
    same("SF1 the year to the final operator: same call", i_cur)
    months = {(r["y"], r["m"]) for r in prod_rows}
    ok("SF2 the production extract holds the rating year's twelve months and every campaign falls inside it",
       months == set(M.RATING_MONTHS) and all(W.RATING_START <= d0 and d1 <= W.RATING_END for _, d0, d1 in W.CAMPAIGNS))
    ok("SF3 'emitted' is the standard's word and the note says a pass does not detect every plume",
       "emitted" in std and "does not detect every plume" in note)
    same("SF4 every detected plume counted, no 5 kg/h floor: same call",
         M.intensity(M.emitted(sites, passes, pl, w, floor=0.0), prod))
    ok("SF5 no producer within 0.01 of the line", min(abs(v - W.LIMIT) for v in inten.values()) >= 0.01,
       "%.3f" % min(abs(v - W.LIMIT) for v in inten.values()))
    same("SF6 Talmora's facility number mapped through the register: same call", main_call(target, attribution="facility")[0])
    same("SF7 a twice-imaged site's first pass: same call", M.intensity(M.emitted(sites, passes, pl, w, first_pass_only=True), prod))
    same("SF7 a twice-imaged site's passes summed: same call", M.intensity(M.emitted(sites, passes, pl, w, sum_passes=True), prod))
    same("SF8 a site on a lost line counted zero: same call", M.intensity(M.emitted(sites, passes, pl, w, coverage_zero=True), prod))
    same("SF9 every production report summed: same call",
         M.intensity(em, M.produced(B["production"], B["composites"], latest_only=False)))
    same("SF9 gas left at 14.65 psia: same call", M.intensity(em, M.produced(B["production"], B["composites"], convert=False)))
    same("SF9 the 2024 composites: same call", M.intensity(em, M.produced(B["production"], B["composites"], year=2024)))
    for label, f in (("probit", M.fit_pod(B["releases"], link="probit")),
                     ("wind added", None), ("complementary log-log", glm_pod(B["releases"], "cloglog")),
                     ("logistic in the rate itself", glm_pod(B["releases"], "logit", x="rate")),
                     ("binned shares", M.binned_pod(B["releases"]))):
        if label == "wind added":
            fw = M.fit_pod(B["releases"], wind=True)
            iv = M.intensity(M.emitted(sites, passes, pl, lambda z: z["rate"] / fw(z["rate"], z["wind"])), prod)
        else:
            iv = M.intensity(M.emitted(sites, passes, pl, lambda z, f=f: z["rate"] / f(z["rate"])), prod)
        same("SF10 the detection curve as %s: same call" % label, iv)
    same("SF11 each campaign a quarter of the year: same call", M.intensity(M.emitted(sites, passes, pl, w, hours="quarter"), prod))
    same("SF11 each campaign its interval between midpoints: same call",
         M.intensity(M.emitted(sites, passes, pl, w, hours="interval"), prod))
    gas_t = M.produced(B["production"], {f: {y: 1.0 for y in v} for f, v in B["composites"].items()})
    i_gas = M.intensity(em, gas_t)
    ok("SF12 over the gas produced instead of its methane would add %s; the standard files methane produced" % names(
        set(M.rated(i_gas)) - gold), "as a share of its methane produced" in std, names(M.rated(i_gas)))
    same("SF13 short tons on both sides: same call", {p: inten[p] for p in W.NAMES})
    # Layer 1 items mishandled (R2)
    same("R2 a one-to-one facility join (each number to one register site): same call",
         main_call(target, passes_mode="one_to_one")[0])
    cst = dt.timezone(dt.timedelta(hours=-6))
    shifted = [q for q in pl if q["site"] and M.operator_on(X["site"][q["site"]], q["time"].replace(tzinfo=dt.timezone.utc)
                                                              .astimezone(cst).date()) != q.get("producer")]
    ok("R2 UTC image times read as Central time: no plume changes producer", not shifted)
    unreg = {r["facility_number"] for r in BD._rows(target, "coverage")} - {f for s in sites for f in s["fac"]}
    near_unreg = [q for q in pl if q["nf"] in unreg and q["site"]]
    ok("R2 the coverage file's unregistered facilities carry no producer plume", not near_unreg, "%d numbers" % len(unreg))

    # ---------------------------------------------------------------------------------------- the asks
    deliv, flar = X["deliv"], X["flar"]
    st_, sn, hv = B["statements"], B["stations"], B["heating"]
    ok("Ask A golden matches figures.json", all(round(deliv[p]) == F["delivered_mmcf"][p] for p in W.NAMES))
    ok("Ask B golden matches figures.json", all(round(flar[p]) == F["flared_mmcf"][p] for p in W.NAMES))
    ea = min(edge(v, 1) for v in list(deliv.values()) + [sum(deliv.values())])
    eb = min(edge(v, 1) for v in list(flar.values()) + [sum(flar.values())])
    ok("Ask A and B every figure and total 0.04 MMcf or more from its rounding edge", ea >= 0.04 and eb >= 0.04,
       "%.3f / %.3f" % (ea, eb))
    stopsA = {"the statement's accounting month (the cracker)": M.delivered(st_, sn, hv, by="statement"),
              "both March revisions": M.delivered(st_, sn, hv, revisions="all")}
    for label, v in stopsA.items():
        moved = [p for p in W.NAMES if round(v[p]) != round(deliv[p])]
        ok("Ask A stop '%s' misses every producer" % label, len(moved) == 12,
           "smallest move %.1f" % min(abs(v[p] - deliv[p]) for p in W.NAMES))
    eo = {sn[s]["op"] for s in sn if "energy meter" in sn[s]["metering"]}
    for label, mode in (("MMBtu read as Mcf", "as_mcf"), ("blank volumes as zero", "drop")):
        v = M.delivered(st_, sn, hv, energy=mode)
        moved = {p for p in W.NAMES if round(v[p]) != round(deliv[p])}
        ok("Ask A stop '%s' moves the two energy-metered stations' producers" % label, moved == eo and len(eo) == 2,
           names(moved))
    edge_ppa = {}
    for ln in st_:
        if ln["kind"] == "ppa" and ((ln["sy"], ln["sm"]), (ln["py"], ln["pm"])) in (((2025, 10), (2025, 9)), ((2026, 10), (2026, 9))):
            edge_ppa.setdefault(sn[ln["station"]]["op"], set()).add(ln["sy"])
    ok("Ask A primary: every producer's October 2025 and October 2026 statements carry a September adjustment",
       all(edge_ppa.get(p) == {2025, 2026} for p in W.NAMES))
    rev2 = {ln["station"] for ln in st_ if ln["rev"] == 2}
    ok("Ask A hazard: the March 2026 statement re-issued whole as revision 2", rev2 == {ln["station"] for ln in st_ if
                                                                                        (ln["sy"], ln["sm"]) == (2026, 3)})
    fl_, le = B["flaring"], B["leases"]
    stopsB = {"FL alone (the cracker)": M.flared(fl_, le, codes=("FL",)),
              "FL, FR and VT": M.flared(fl_, le, codes=("FL", "FR", "VT")),
              "every report version": M.flared(fl_, le, versions="all")}
    for label, v in stopsB.items():
        moved = [p for p in W.NAMES if round(v[p]) != round(flar[p])]
        ok("Ask B stop '%s' misses every producer" % label, len(moved) == 12,
           "smallest move %.1f" % min(abs(v[p] - flar[p]) for p in W.NAMES))
    ve = M.flared(fl_, le, at_end=True)
    moved = {p for p in W.NAMES if round(ve[p]) != round(flar[p])}
    ok("Ask B stop 'each lease to its final operator' moves the transfers' four producers", moved == {
        "Calbrent Energy", "Sunmere Natural Gas", "Ruskmoor Exploration", "Pellisto Gas Partners"}, names(moved))
    lat = M.latest_flaring(fl_)
    bs = B["basin"]
    tie = all(bs[k]["flared"] == sum(r["mcf"] for r in lat if (r["y"], r["m"]) == k and r["code"] in ("FL", "FR"))
              for k in bs)
    fl_only = [k for k in bs if bs[k]["flared"] == sum(r["mcf"] for r in lat if (r["y"], r["m"]) == k and r["code"] == "FL")]
    ok("Ask B referee: the basin summary ties with FL and FR every month, with FL alone only before 2026",
       tie and all(k < W.CODE_SWITCH for k in fl_only) and len(fl_only) == sum(1 for k in bs if k < W.CODE_SWITCH))
    ok("Ask B primary: FR appears only from January 2026", min((r["y"], r["m"]) for r in fl_ if r["code"] == "FR") == W.CODE_SWITCH)

    # the separation line and the distractors
    main_files = {W.FILES[k] for k in W.MAIN_FILES}
    tmp = tempfile.mkdtemp(dir=TASK, suffix=".chk")
    try:
        for f in files:
            if f not in {W.FILES[k] for k in ("statements", "stations", "heating", "flaring", "leases", "basin")}:
                shutil.copy(os.path.join(target, f), tmp)
        i2 = main_call(tmp)[0]
        ok("separation: the call computed without the plant and flaring files is identical",
           all(abs(i2[p] - inten[p]) < 1e-12 for p in W.NAMES))
        shutil.rmtree(tmp)
        os.makedirs(tmp)
        for f in files:
            if f not in main_files - {W.FILES["operators"]}:
                shutil.copy(os.path.join(target, f), tmp)
        d2, f2 = asks(tmp)
        ok("separation: the asks computed without the plume, coverage, release, register, production and gas files are "
           "identical", all(abs(d2[p] - deliv[p]) < 1e-9 and abs(f2[p] - flar[p]) < 1e-9 for p in W.NAMES))
        for k in W.DISTRACTORS:
            shutil.rmtree(tmp)
            os.makedirs(tmp)
            for f in files:
                if f != W.FILES[k]:
                    shutil.copy(os.path.join(target, f), tmp)
            i3 = main_call(tmp)[0]
            d3, f3 = asks(tmp)
            ok("distractor %s: the call and every ask unchanged without it" % W.FILES[k],
               all(abs(i3[p] - inten[p]) < 1e-12 and abs(d3[p] - deliv[p]) < 1e-9 and abs(f3[p] - flar[p]) < 1e-9
                   for p in W.NAMES))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    ask_rows = sum(1 for _ in open(os.path.join(target, W.FILES["statements"]))) + sum(
        1 for _ in open(os.path.join(target, W.FILES["flaring"])))
    ok("separation: no device or hazard row sits in a file the call reads (the devices live in %d plant and flaring "
       "rows)" % ask_rows, not ({W.FILES[k] for k in ("statements", "flaring", "leases", "basin", "stations", "heating")}
                                & main_files))

    # ------------------------------------------------------------------------------- input gates and hygiene
    fmts = {f.rsplit(".", 1)[1] for f in files}
    rows = {f: sum(1 for _ in open(os.path.join(target, f), encoding="utf-8")) - 1 for f in files if f.endswith(".csv")}
    ok("input gates: 10 or more files, 3 or more formats, a file of 25,000 rows or more",
       len(files) >= 10 and len(fmts) >= 3 and max(rows.values()) >= 25000, "%d files, %s, largest %d rows" % (
           len(files), sorted(fmts), max(rows.values())))
    ok("two distractors named in metadata.json, none in a file name",
       len(meta["distractor_files"]) == 2 and all(f in files for f in meta["distractor_files"]) and
       not any("distract" in f for f in files))
    man = [r.split(",")[0] for r in open(os.path.join(target, W.FILES["manifest"]), encoding="utf-8").read().splitlines()[1:]]
    ok("the manifest lists every shipped file but itself", set(man) == set(files) - {W.FILES["manifest"]})
    dic = openpyxl.load_workbook(os.path.join(target, W.FILES["dict"]))["Field definitions"]
    named = {c.value for c in dic["A"] if c.value}
    dic_text = " ".join(str(c.value) for row in dic.iter_rows() for c in row if c.value)
    ok("the dictionary names every data file in the folder", {f for f in files if not f.endswith(".pdf")} - {
        W.FILES["dict"]} <= named)
    ok("the dictionary's disposition codes are the commission's 2024 pair (FL flared, VT vented)",
       "FL flared, VT vented" in dic_text and " FR " not in dic_text)
    texts = [open(os.path.join(target, f), encoding="utf-8").read() for f in files if f.endswith((".csv", ".json"))]
    texts += list(doc_text.values()) + [dic_text]
    texts += [open(os.path.join(TASK, n), encoding="utf-8").read() for n in ("prompt.txt", "submission.md", "DESIGN_NOTE.md")]
    ok("no em or en dash in the pack, the prompt, the submission or the design note",
       not any("—" in t or "–" in t for t in texts))
    ok("the only people named are the drawn personas", W.PEOPLE["ratings"] in std and W.PEOPLE["provider"] in note and
       not re.search(r"John Doe|Jane Doe|Acme", alltext))
    pull = W.PACK_DATE - dt.timedelta(days=3)
    ok("pull dates hold: no report received after the commission's pull, no plume outside its campaign, no statement "
       "after October 2026", max(r["filed"] for r in prod_rows) <= pull and max(r["filed"] for r in fl_) <= pull and
       all(d0 <= q["time"].date() <= d1 for c, d0, d1 in W.CAMPAIGNS for q in pl if q["campaign"] == c) and
       max((ln["sy"], ln["sm"]) for ln in st_) == (2026, 10))
    reg = {}
    for s in sites:
        reg[s["ops"][-1][1]] = reg.get(s["ops"][-1][1], 0) + 1
    ok("realism: the producers' site counts are distinct and none a multiple of five",
       len(set(reg.values())) == 12 and all(v % 5 for v in reg.values()), str(sorted(reg.values())))
    gas_bcf = {}
    for r in prod_rows:
        if r["kind"] == "ORIGINAL" and r["op"] in W.NAMES:
            gas_bcf[r["op"]] = gas_bcf.get(r["op"], 0) + r["gas"] / 1e6
    ok("realism: no producer's year of original reports a whole number of Bcf to within 1,000 Mcf", all(
        abs(v - round(v)) > 0.001 for v in gas_bcf.values()), ", ".join("%.3f" % v for v in gas_bcf.values()))
    per_camp = [sum(1 for q in pl if q["campaign"] == c) for c, _, _ in W.CAMPAIGNS]
    ok("realism: each campaign's plume count differs", len(set(per_camp)) == 4, str(per_camp))
    by_line = {}
    for q in pl:
        by_line.setdefault((q["campaign"], q["line"]), set()).add(q["wind"])
    rat = np.array([q["sigma"] / q["rate"] for q in pl if q["rate"] >= 10])
    ok("realism: the reanalysis wind varies along each flight line and the uncertainty is no fixed share of the rate",
       all(len(v) >= 3 for k, v in by_line.items() if sum(1 for q in pl if (q["campaign"], q["line"]) == k) >= 20) and
       rat.std() > 0.02, "%d winds, uncertainty share sd %.3f" % (len({q["wind"] for q in pl}), rat.std()))
    for f in files:
        if f.endswith(".pdf"):
            raw = open(os.path.join(target, f), "rb").read()
            ok("container: %s carries no writer name" % f, b"ReportLab" not in raw and b"reportlab" not in raw)
        if f.endswith(".xlsx"):
            with zipfile.ZipFile(os.path.join(target, f)) as z:
                core = z.read("docProps/core.xml").decode()
            ok("container: %s stamped with the fiction's people and dates" % f, "openpyxl" not in core and
               ("2025-11" in core or "2026-11" in core))

    # ----------------------------------------------------------------------------------------- the golden files
    import csv
    out = list(csv.DictReader(open(os.path.join(golden, W.GOLDEN["plumes"]), encoding="utf-8")))
    ok("the conformed file: one row per plume, ids unique, %d with a producer site" % F["plumes_attributed"],
       len(out) == len(pl) == F["csv_rows"] and len({r["plume_id"] for r in out}) == len(out) and
       sum(1 for r in out if r["site_id"]) == F["plumes_attributed"] and all(
           (r["site_id"] == "") == (r["producer"] == "") for r in out))
    wb = openpyxl.load_workbook(os.path.join(golden, W.GOLDEN["workbook"]))
    ok("the workbook's sheets", wb.sheetnames == ["2027 ratings", "Intensity chart", "Detection curve", "Sources"])
    ok("the workbook's chart", len(wb["Intensity chart"]._charts) == 1)
    ws = wb["2027 ratings"]
    got = {ws.cell(row=r, column=1).value: [ws.cell(row=r, column=c).value for c in range(2, 8)] for r in range(5, 17)}
    ok("the workbook carries every producer's figures as figures.json has them", all(
        got[p][1] == F["emitted_t"][p] and abs(got[p][2] - F["intensity"][p]) < 1e-9 and
        got[p][3] == ("Low-methane" if p in gold else "Not rated") and got[p][4] == F["delivered_mmcf"][p] and
        got[p][5] == F["flared_mmcf"][p] for p in W.NAMES))
    ok("figures.json carries the rated four", set(F["rated"]) == gold)
    ee = min(edge(v, 1) for v in em.values())
    ei = min(edge(v, 0.001) for v in inten.values())
    ok("each producer's methane emitted 0.05 t or more from its rounding edge, each intensity 0.00005 or more",
       ee >= 0.05 and ei >= 0.00005, "%.3f t / %.6f" % (ee, ei))
    return report()


def report():
    bad = [r for r in RESULTS if not r[1]]
    for label, good, detail in RESULTS:
        print("%s  %s%s" % ("PASS" if good else "FAIL", label, ("  [" + detail + "]") if detail else ""))
    print("checks: %d run, %d failed" % (len(RESULTS), len(bad)))
    return not bad


if __name__ == "__main__":
    good = run(os.path.join(TASK, "target"), os.path.join(TASK, "golden"))
    raise SystemExit(0 if good else 1)
