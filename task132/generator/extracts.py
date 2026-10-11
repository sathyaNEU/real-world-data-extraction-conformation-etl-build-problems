"""The system extracts: the EAM asset register, the PSU rotating-item history, the work-order export,
substations, spares stock and Veldtek's RMA register."""
import datetime as dt

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font

import crewtext
import sim as S
import world as W

MODELS = {
    "VPR-340": ("Veldtek Protection Systems", "Feeder protection"),
    "VPR-220": ("Veldtek Protection Systems", "Feeder protection"),
    "NX-9": ("Norland Controls", "Feeder protection"),
    "AT-7": ("Arcwell Automation", "Recloser control"),
    "NX-5": ("Norland Controls", "Transformer differential"),
    "KDR-30": ("Kestrel Grid Devices", "Bus protection"),
    "RT-400": ("Arcwell Automation", "Substation RTU"),
    "RT-600": ("Arcwell Automation", "Substation RTU"),
}
SERIAL = {"VPR-340": "V340-%02d-%05d", "VPR-220": "V220-%02d-%05d", "NX-9": "N9-%02d-%05d", "NX-5": "N5-%02d-%05d",
          "AT-7": "AW7%02d%05d", "KDR-30": "K30%02d%04d", "RT-400": "AW4%02d%05d", "RT-600": "AW6%02d%05d"}
GOLIVE = W.MIGRATION
END = dt.date(2026, 10, 1)
# the device that goes into a bay when the old one comes out
NEW_DEVICE = {"VPR-340": "NX-9", "VPR-220": "NX-9", "NX-9": "NX-9", "NX-5": "NX-5", "KDR-30": "KDR-30",
              "RT-400": "RT-600", "AT-7": "AT-7"}
MFG_LAG = (20, 270)                    # days from manufacture to installation of a relay, control or RTU
HELD_LAG = (365, 1096)                 # the same for one held in stores for a year or more
STORES = {"North": "North Depot", "Central": "Central Stores (Barrie Road)", "South": "South Depot"}
YEAR_ONLY_BEFORE = dt.date(2001, 1, 1)  # legacy records older than this sometimes held only the year
RMA_START = dt.date(2015, 4, 1)         # Veldtek's register for Tarrock's VPR-340 PSU returns opens with SB-15-02
REG_SAVED = dt.date(2026, 9, 29)        # Veldtek's last save of the register before the count
CUTOVER_HOLD = dt.date(2020, 12, 14)    # stores held shipments over the EAM cutover


def _d(rng, y0, y1):
    """a working day in a month drawn from y0 to y1"""
    first = dt.date(int(rng.integers(y0, y1 + 1)), int(rng.integers(1, 13)), 1)
    return S.pick_bday(rng, first, S.month_add(first, 1) - dt.timedelta(days=1))


def bay_plan(fl, prog):
    """a few stations keep older VPR-220 feeders beside their VPR-340s, and a few have a spare feeder bay after them"""
    rx = np.random.default_rng([W.SEED, 40])
    n = fl.groupby("sub_code").size()
    ok = sorted(s for s, r in zip(prog["sub_code"], prog["retired_on"]) if not isinstance(r, dt.date) and n[s] >= 6)
    pick = [str(x) for x in rx.choice(ok, 9, replace=False)]
    plan = {s: dict(vpr220=int(rx.integers(2, 5)), spare=0) for s in pick[:5]}
    plan.update({s: dict(vpr220=0, spare=1) for s in pick[5:]})
    return plan


def assets(rng, subs, fl, prog, plan):
    """every protection asset at every station since 1988, with its bay, install and retirement dates"""
    cut = {sc: d for sc, d in zip(prog["sub_code"], prog["cut_days"]) if isinstance(d, list)}
    camp = fl.groupby("sub_code")["install"].min().to_dict()
    rv = np.random.default_rng([W.SEED, 43])
    out = []
    for _, s in subs.iterrows():
        sc = s.sub_code
        ts = s.station_type == "TS"
        days = cut.get(sc)
        mine = fl[fl["sub_code"] == sc]
        for _, r in mine.iterrows():
            out.append(dict(model="VPR-340", sub_code=sc, bay=r.bay, install=r.install,
                            retired=r.retire if isinstance(r.retire, dt.date) else None, relay_key=r.relay_key))
        nv = len(mine)
        if nv:
            bp = plan.get(sc, dict(vpr220=0, spare=0))
            # older VPR-220 feeders kept beside the VPR-340s, then a spare bay, then the newer feeders
            y0 = camp[sc].year
            for j in range(bp["vpr220"]):
                out.append(dict(model="VPR-220", sub_code=sc, bay="F%02d" % (nv + j + 1),
                                install=_d(rv, max(1989, min(y0 - 10, 2000)), min(2004, max(y0 - 1, 1993))),
                                retired=None, relay_key=None))
            off = nv + bp["vpr220"] + bp["spare"]
            extra = int(rng.integers(1, 5)) if ts else int(rng.integers(0, 3))
            for j in range(extra):
                y0n = max(camp[sc].year + 1, 2011)
                out.append(dict(model="NX-9", sub_code=sc, bay="F%02d" % (off + j + 1), install=_d(rng, min(y0n, 2026), 2026),
                                retired=None, relay_key=None))
        else:
            nfeed = int(rng.integers(12, 21)) if ts else int(rng.integers(5, 13))
            rural = (not ts) and rng.random() < 0.35
            for j in range(nfeed):
                bay = "F%02d" % (j + 1)
                u = rng.random()
                if rural and u < 0.45:
                    out.append(dict(model="AT-7", sub_code=sc, bay=bay, install=_d(rng, 2000, 2020), retired=None, relay_key=None))
                elif u < 0.62:
                    inst = _d(rng, 1989, 2004)
                    ret = None
                    if rng.random() < 0.13:
                        ret = _d(rng, 2021, 2026)
                        if ret > W.ASOF:
                            ret = None
                    out.append(dict(model="VPR-220", sub_code=sc, bay=bay, install=inst, retired=ret, relay_key=None))
                    if ret:
                        out.append(dict(model="NX-9", sub_code=sc, bay=bay, install=ret, retired=None, relay_key=None))
                else:
                    out.append(dict(model="NX-9", sub_code=sc, bay=bay, install=_d(rng, 2010, 2026), retired=None, relay_key=None))
        for t in range(1, s.n_tx + 1):
            out.append(dict(model="NX-5", sub_code=sc, bay="T%d" % t, install=_d(rng, 1995, 2024), retired=None, relay_key=None))
        for b in range(1, s.n_bus + 1):
            out.append(dict(model="KDR-30", sub_code=sc, bay="B%d" % b, install=_d(rng, 1996, 2022), retired=None, relay_key=None))
        rtu = _d(rng, 2003, 2016)
        rret = None
        if days is None and rng.random() < 0.05:
            rret = _d(rng, 2021, 2025)
            out.append(dict(model="RT-600", sub_code=sc, bay="RTU", install=rret, retired=None, relay_key=None))
        out.append(dict(model="RT-400", sub_code=sc, bay="RTU", install=rtu, retired=rret, relay_key=None))
        if days:
            # station rebuild: each bay is cut over on its own working day and the new device goes in the day the old
            # one comes out; a device fitted in the two years before the rebuild stays in service
            keep_from = days[0] - dt.timedelta(days=730)
            new = []
            for o in [o for o in out if o["sub_code"] == sc]:
                if o["model"] == "VPR-340":
                    new.append(dict(model="NX-9", sub_code=sc, bay=o["bay"], install=o["retired"], retired=None, relay_key=None))
                elif o["retired"] is None and o["install"] < keep_from:
                    o["retired"] = S.cut_date(days, o["bay"], nv)
                    new.append(dict(model=NEW_DEVICE[o["model"]], sub_code=sc, bay=o["bay"], install=o["retired"],
                                    retired=None, relay_key=None))
            out.extend(new)
    a = pd.DataFrame(out)
    a = a[(a["retired"].isna() | (a["retired"] >= GOLIVE)) & (a["install"] <= W.ASOF)].copy()
    return a.reset_index(drop=True)


def register(rng, subs, fl, prog, plan):
    """protection asset register: every asset in service on 1 Jan 2021 or installed since. Returns the register, the
    asset number of each VPR-340 relay, each VPR-340 relay's date of manufacture, and the VPR-340 relays whose legacy
    record held only the year of installation."""
    a = assets(rng, subs, fl, prog, plan)
    leg = dict(zip(subs["sub_code"], subs["legacy_region"]))
    cur = dict(zip(subs["sub_code"], subs["region"]))
    a["loaded"] = a["install"] < GOLIVE
    # EAM asset numbers: the go-live load in location order, then new assets as they were created; other asset
    # classes share the sequence, so protection numbers have gaps
    a["_o"] = [(0, s, b, str(i)) if l else (1, str(i), s, b) for s, b, i, l in zip(a["sub_code"], a["bay"], a["install"], a["loaded"])]
    a = a.sort_values("_o", key=lambda c: c.map(lambda t: t), kind="mergesort").reset_index(drop=True)
    num, ids = 1000000, []
    for l in a["loaded"]:
        num += int(rng.integers(1, 9)) if l else int(rng.integers(4, 61))
        if not l and num < 1046000:
            num = 1046000 + int(rng.integers(0, 400))
        ids.append("P%07d" % num)
    a["asset_id"] = ids
    # manufacturer serials: the year of manufacture and the maker's running number, in order of manufacture (the
    # maker numbers every unit it builds, so ours have gaps); a unit that sat in stores goes in a year or more later
    lags = []
    for i in a["install"]:
        held = rng.random() < 0.13
        lags.append(int(rng.integers(*HELD_LAG)) if held else int(rng.integers(*MFG_LAG)))
    a["mfg"] = [i - dt.timedelta(days=g) for i, g in zip(a["install"], lags)]
    a["serial_no"] = ""
    for m, g in a.groupby("model"):
        seq = int(rng.integers(100, 3000))
        for i in g.sort_values(["mfg", "asset_id"], kind="mergesort").index:
            seq += int(rng.integers(1, 23))
            a.at[i, "serial_no"] = SERIAL[m] % (a.at[i, "mfg"].year % 100, seq)
    a["manufacturer"] = [MODELS[m][0] for m in a["model"]]
    a["function"] = [MODELS[m][1] for m in a["model"]]
    a["status"] = ["RETIRED" if r is not None and not pd.isna(r) else "IN SERVICE" for r in a["retired"]]
    a["region"] = [leg[s] if (l or i < W.REORG) else cur[s] for s, l, i in zip(a["sub_code"], a["loaded"], a["install"])]
    # legacy lists from before the legacy system often held only the year a device went in; those were loaded as
    # 1 July of that year (whole stations' records were kept that way)
    ry = np.random.default_rng([W.SEED, 41])
    yearly_station = {sc: ry.random() < 0.4 for sc in sorted(subs["sub_code"])}
    yo = [bool(l and i < YEAR_ONLY_BEFORE and ry.random() < (0.8 if yearly_station[s] else 0.05))
          for s, l, i in zip(a["sub_code"], a["loaded"], a["install"])]
    a["install_date"] = [dt.date(i.year, 7, 1) if y else i for i, y in zip(a["install"], yo)]
    a["retired_date"] = [None if (r is None or pd.isna(r)) else r for r in a["retired"]]
    a["_yo"] = yo
    reg = a.sort_values(["sub_code", "bay", "install_date", "asset_id"], kind="mergesort").reset_index(drop=True)
    cols = ["asset_id", "model", "manufacturer", "function", "sub_code", "bay", "serial_no", "install_date", "status",
            "retired_date", "region"]
    keyed = [(int(k), aid, mf, y) for k, aid, mf, y in zip(reg["relay_key"], reg["asset_id"], reg["mfg"], reg["_yo"])
             if k is not None and not pd.isna(k)]
    return (reg[cols], {k: aid for k, aid, _, _ in keyed}, {k: mf for k, _, mf, _ in keyed},
            {k for k, _, _, y in keyed if y})


# ---------------------------------------------------------------- station equipment the P&C crews work on

def equipment(subs, reg):
    """station DC systems (battery banks and chargers), the breaker in each relay-protected bay and the bus VTs: the
    assets the P&C crews test beside the relays. Their EAM numbers come from the same sequence as the protection
    assets: the go-live load numbered each station's equipment in location order, newer equipment as it was created."""
    re_ = np.random.default_rng([W.SEED, 44])
    rows = []
    prot = reg[reg["model"].isin(["VPR-340", "VPR-220", "NX-9", "NX-5", "KDR-30"])]
    first = prot.groupby(["sub_code", "bay"])["install_date"].min().to_dict()
    for s in subs.itertuples():
        ts = s.station_type == "TS"
        for kind, names, y0, y1, life in (("battery", ["BAT1", "BAT2"] if ts else ["BAT1"], 2004, 2019, 0.65),
                                          ("charger", ["CHG1", "CHG2"] if ts else ["CHG1"], 2002, 2019, 0.15)):
            for nm in names:
                inst = _d(re_, y0, y1)
                old = inst < dt.date(2009, 1, 1) if kind == "battery" else True
                ret = _d(re_, 2021, 2026) if old and re_.random() < life else None
                if ret is not None and ret > W.ASOF:
                    ret = None
                rows.append(dict(cls=kind, sub_code=s.sub_code, name=nm, install=inst, retired=ret))
                if ret is not None:
                    rows.append(dict(cls=kind, sub_code=s.sub_code, name=nm, install=ret, retired=None))
        for (sc, bay), f0 in sorted(first.items()):
            if sc != s.sub_code or bay.startswith("B"):
                continue
            inst = f0 if f0 >= GOLIVE else _d(re_, 1972, max(1973, min(f0.year, 2016)))
            rows.append(dict(cls="cb", sub_code=sc, name=bay, install=inst, retired=None))
        for b in range(1, s.n_bus + 1):
            rows.append(dict(cls="vt", sub_code=s.sub_code, name="VT-B%d" % b, install=_d(re_, 1980, 2015), retired=None))
    eq = pd.DataFrame(rows)
    eq["asset_id"] = _alloc_ids(reg, eq, re_)
    return eq


def _alloc_ids(reg, eq, rng):
    used = {int(a[1:]) for a in reg["asset_id"]}
    out = [None] * len(eq)
    order = {"B": 0, "F": 1, "T": 2, "V": 3, "C": 4}
    loaded = reg[[d < GOLIVE for d in reg["install_date"]]]
    blocks = {sc: sorted(int(a[1:]) for a in g["asset_id"]) for sc, g in loaded.groupby("sub_code")}
    for sc, g in eq[[d < GOLIVE for d in eq["install"]]].groupby("sub_code"):
        items = sorted(g.index, key=lambda i: (order.get(eq.at[i, "name"][0], 5), eq.at[i, "name"], eq.at[i, "cls"]))
        nums = blocks.get(sc) or [1045000]
        free = [n for n in range(nums[0] + 1, nums[-1]) if n not in used]
        k = nums[-1] + 1
        while len(free) < len(items):
            if k not in used:
                free.append(k)
            k += 1
        step = len(free) / max(1, len(items))
        for j, i in enumerate(items):
            n = free[int(j * step)]
            used.add(n)
            out[i] = "P%07d" % n
    new = sorted((d, int(a[1:])) for d, a in zip(reg["install_date"], reg["asset_id"]) if d >= GOLIVE)
    for i in sorted(eq.index[[d >= GOLIVE for d in eq["install"]]], key=lambda i: (eq.at[i, "install"], i)):
        d = eq.at[i, "install"]
        lo = max([n for dd, n in new if dd <= d] or [1046000])
        hi = min([n for dd, n in new if dd > d] or [lo + 400])
        cand = [n for n in range(lo + 1, hi) if n not in used] or [n for n in range(hi + 1, hi + 400) if n not in used]
        n = cand[int(rng.integers(len(cand)))]
        used.add(n)
        out[i] = "P%07d" % n
    return out


# ---------------------------------------------------------------- the power supply modules and their returns

def legacy_wo(d, rng):
    """a work-order number from the legacy maintenance system (shared with the lines and metering crews)"""
    return "%d" % (468000 + int((d - dt.date(2015, 1, 1)).days * 13.6) + int(rng.integers(0, 40)))


def psu_cycle(sp, f, key2aid, relay_mfg, sub_of, depot_of):
    """Module serials for every spell; failed modules, and the modules the 2015 service bulletin took out, sent back
    to Veldtek under RMA once its register was running, then repaired, found fault-free or scrapped; repaired modules
    come back to stores and some are fitted again. A new module's serial carries its year of manufacture and Veldtek's
    running number in order of manufacture; a repaired module keeps its serial."""
    sp = sp.copy()
    r_ship = np.random.default_rng([W.SEED, 21, 1])
    r_res = np.random.default_rng([W.SEED, 21, 2])
    r_pool = np.random.default_rng([W.SEED, 21, 3])
    r_made = np.random.default_rng([W.SEED, 21, 4])
    r_ref = np.random.default_rng([W.SEED, 21, 5])
    spell_of = {(int(k), int(m)): i for i, k, m in zip(sp.index, sp["relay_key"], sp["module"])}
    cand = []
    for r in f.sort_values(["done", "relay_key"], kind="mergesort").itertuples():
        if r.done >= RMA_START:
            cand.append(dict(kind="EAM" if r.month >= GOLIVE else "LEGACY", relay_key=int(r.relay_key), module=int(r.module),
                             removed=r.done, reported=r.reported))
    camp = sp[sp["reason"] == "CAMPAIGN"]
    for k, m, d in zip(camp["relay_key"], camp["module"], camp["removed"]):
        cand.append(dict(kind="SB", relay_key=int(k), module=int(m), removed=d, reported=d))
    cand.sort(key=lambda c: (c["removed"], c["relay_key"], c["module"]))
    ships, batch = [], {}
    for c in cand:
        if c["kind"] == "SB":
            if r_ship.random() > 0.86:
                continue
            sub = sub_of[c["relay_key"]]
            if sub not in batch:
                batch[sub] = (len(batch), S.add_bdays(c["removed"], int(r_ship.integers(12, 40))))
            g, ship = batch[sub]
            ships.append(dict(c, ship=ship, group="SB%d" % g, ref=None))
            continue
        if r_ship.random() > (0.5 if c["kind"] == "LEGACY" else 0.55):
            continue
        u = r_ship.random()
        lag = (int(r_ship.integers(2, 9)) if u < 0.55 else int(r_ship.integers(9, 19)) if u < 0.82 else
               int(r_ship.integers(19, 46)) if u < 0.95 else int(r_ship.integers(46, 111)))
        ship = S.add_bdays(c["removed"], lag)
        if c["kind"] == "LEGACY" and ship >= CUTOVER_HOLD:
            ship = S.add_bdays(dt.date(2021, 1, 4), int(r_ship.integers(0, 9)))
        ships.append(dict(c, ship=ship, group="M%d" % len(ships), ref=None))
    # failed modules waiting at the same depot within a couple of weeks go to Veldtek in one box
    eam = sorted([s for s in ships if s["kind"] == "EAM"], key=lambda s: (s["ship"], s["relay_key"]))
    paired = set()
    for i, a in enumerate(eam):
        if id(a) in paired:
            continue
        for b in eam[i + 1:]:
            if (b["ship"] - a["ship"]).days > 14:
                break
            if id(b) in paired or depot_of[b["relay_key"]] != depot_of[a["relay_key"]]:
                continue
            if r_ship.random() < 0.5:
                a["ship"] = b["ship"]
                b["group"] = a["group"]
                paired.update((id(a), id(b)))
            break
    groups = {}
    for s in ships:
        groups.setdefault(s["group"], []).append(s)
    keep = []
    for g in sorted(groups, key=lambda g: (min(s["ship"] for s in groups[g]), g)):
        members = groups[g]
        ship = max(s["ship"] for s in members)
        lag = int(r_ship.integers(1, 4)) if r_ship.random() < 0.85 else int(r_ship.integers(4, 10))
        recv = S.add_bdays(ship, lag)
        if recv > REG_SAVED:
            continue                                       # still on the shelf at the depot at the count
        for s in members:
            s["ship"], s["received"] = ship, recv
            keep.append(s)
    for s in keep:
        if s["kind"] == "SB":
            s["result"], s["returned"] = "SCRAPPED (SB-15-02)", None
            continue
        evald = S.add_bdays(s["received"], int(r_res.integers(4, 13)))
        u = r_res.random()
        res = "REPAIRED" if u < 0.77 else ("NO FAULT FOUND" if u < 0.82 else "BEYOND ECONOMIC REPAIR")
        if res == "NO FAULT FOUND":
            t = int(r_res.integers(5, 19))
        else:
            v = r_res.random()
            t = int(r_res.integers(14, 33)) if v < 0.7 else int(r_res.integers(32, 61)) if v < 0.92 else int(r_res.integers(60, 121))
            if s["received"].year >= 2026:
                t = int(t * 1.25)
        back = None if res == "BEYOND ECONOMIC REPAIR" else S.add_bdays(s["received"], t)
        if evald > REG_SAVED or (back is not None and back > REG_SAVED):
            res, back = "OPEN", None
        s["result"], s["returned"] = res, back
    # repaired and fault-free modules back on the stores shelf a few days after Veldtek ships them
    shelf = sorted((S.add_bdays(s["returned"], int(r_pool.integers(1, 6))), spell_of[(s["relay_key"], s["module"])])
                   for s in keep if s["returned"] is not None)
    taken = set()
    made, origin = {}, {}
    for i in sp.sort_values(["fitted", "relay_key", "module"], kind="mergesort").index:
        when = sp.at[i, "fitted"]
        key, mod = int(sp.at[i, "relay_key"]), int(sp.at[i, "module"])
        if mod == 0:
            made[i] = relay_mfg[key] - dt.timedelta(days=int(r_made.integers(7, 45)))     # built into the relay
        elif sp.at[i - 1, "reason"] == "CAMPAIGN":
            made[i] = when - dt.timedelta(days=int(r_made.integers(30, 160)))            # built for the bulletin
        else:
            avail = [p for p in shelf if p[0] <= when and p[1] not in taken and p[1] != i]
            if avail and r_pool.random() < 0.45:
                pick = avail[0]
                taken.add(pick[1])
                origin[i] = pick[1]
            else:
                made[i] = when - dt.timedelta(days=int(r_made.integers(45, 900)))       # from stores stock
    serial, n = {}, 21000
    for i in sorted(made, key=lambda j: (made[j], j)):
        n += int(r_made.integers(1, 9))
        serial[i] = "PS02-%02d-%05d" % (made[i].year % 100, n)

    def serial_of(i):
        while i in origin:
            i = origin[i]
        return serial[i]

    sp["serial"] = [serial_of(i) for i in sp.index]
    sp["repaired"] = [i in origin for i in sp.index]
    sp["asset_id"] = [key2aid[int(k)] for k in sp["relay_key"]]
    ships = pd.DataFrame(keep)
    ships["spell"] = [spell_of[(k, m)] for k, m in zip(ships["relay_key"], ships["module"])]
    ships["serial"] = [serial_of(i) for i in ships["spell"]]
    # Tarrock's reference on each return: the EAM logistics work order (filled in later), or the legacy work order
    sbref = {}
    refs = []
    for s in ships.itertuples():
        if s.kind == "SB":
            if s.group not in sbref:
                sbref[s.group] = legacy_wo(s.removed, r_ref)
            refs.append(sbref[s.group])
        elif s.kind == "LEGACY":
            refs.append(legacy_wo(s.reported, r_ref))
        else:
            refs.append(None)
    ships["ref"] = refs
    ships = ships.sort_values(["ship", "group", "removed", "relay_key"], kind="mergesort").reset_index(drop=True)
    return sp, ships


def refit_plan(f, sp):
    """after some failures the spare goes in on a follow-up work order the same or the next working day (no spare on
    the truck at a call-out, or a day crew finishing what the night crew made safe); the module is fitted no later
    than the first day of the month after the removal, so it counts in service from the same month either way"""
    rp = np.random.default_rng([W.SEED, 23])
    out = {}
    for r in f[f["month"] >= GOLIVE].sort_values(["reported", "relay_key"], kind="mergesort").itertuples():
        if rp.random() > (0.35 if r.em else 0.14):
            continue
        done = r.done
        limit = min(S.month_add(dt.date(done.year, done.month, 1), 1), W.ASOF)
        raised = done if rp.random() < 0.55 else S.add_bdays(done + dt.timedelta(days=1), 0)
        fitted = raised if rp.random() < 0.85 else S.add_bdays(raised + dt.timedelta(days=1), 0)
        if fitted > limit:
            raised = fitted = done
        out[(int(r.relay_key), int(r.module))] = (raised, fitted)
    return out


# ---------------------------------------------------------------- the work-order export

CREWS_LEGACY = {"Lakeshore": [("Lakeshore P&C 1", 0.6), ("Lakeshore P&C 2", 0.4)], "Ridge": [("Ridge P&C", 1.0)],
                "Valley": [("Valley P&C 1", 0.55), ("Valley P&C 2", 0.45)], "Harbour": [("Harbour P&C", 0.8), ("Harbour Test", 0.2)]}
CREWS_CURRENT = {"North": [("North P&C 1", 0.57), ("North P&C 2", 0.43)],
                 "Central": [("Central P&C 1", 0.42), ("Central P&C 2", 0.36), ("Central P&C 3", 0.22)],
                 "South": [("South P&C 1", 0.52), ("South P&C 2", 0.48)]}
# hours a crew books for the round trip to site, and its pace on the job; the north and lakeshore crews cover the
# rural stations
CREW = {"Lakeshore P&C 1": (1.2, 1.00), "Lakeshore P&C 2": (1.4, 1.05), "Ridge P&C": (0.9, 1.00),
        "Valley P&C 1": (0.7, 0.95), "Valley P&C 2": (0.8, 1.05), "Harbour P&C": (0.5, 1.00), "Harbour Test": (0.4, 0.90),
        "North P&C 1": (1.5, 1.00), "North P&C 2": (1.7, 1.05), "Central P&C 1": (0.4, 0.95), "Central P&C 2": (0.5, 1.00),
        "Central P&C 3": (0.6, 1.10), "South P&C 1": (0.8, 1.00), "South P&C 2": (0.9, 0.95)}
KIND = {"VPR-340": "relay", "VPR-220": "relay", "NX-9": "relay", "NX-5": "relay", "KDR-30": "relay", "AT-7": "recloser",
        "RT-400": "rtu", "RT-600": "rtu"}
# corrective work orders per device-year at ten years old; older devices fail more (rate x (0.6 + 0.04 x age))
CM_BASE = {"VPR-340": 0.035, "VPR-220": 0.05, "NX-9": 0.035, "NX-5": 0.04, "KDR-30": 0.035, "AT-7": 0.06,
           "RT-400": 0.055, "RT-600": 0.04}
EQ_CM = {"battery": 0.03, "charger": 0.05, "cb": 0.012, "vt": 0.01}
COMP = {"relay": (["CPU", "IO", "DISP", "COMMS", "PSU"], [0.18, 0.33, 0.12, 0.22, 0.15]),
        "recloser": (["PSU", "CPU", "IO", "DISP", "COMMS"], [0.35, 0.12, 0.25, 0.08, 0.20]),
        "rtu": (["CPU", "IO", "COMMS", "PSU"], [0.15, 0.35, 0.35, 0.15])}
HOLD = {"relay": "backup protection in service", "rtu": "affected points on manual reporting",
        "recloser": "recloser left on non-reclose"}
PART = {"CPU": "processor card", "IO": "IO card", "DISP": "display", "COMMS": "comms card", "PSU": "power supply"}
MAKER = {"Veldtek Protection Systems": "Veldtek", "Norland Controls": "Norland", "Arcwell Automation": "Arcwell",
         "Kestrel Grid Devices": "Kestrel"}
# P&C Standard 3.2 test intervals (days): relays two-yearly at transformer stations and three-yearly at distribution
# stations, recloser controls and RTUs yearly; station DC, breaker trip and instrument-transformer checks
PM_DAYS = {"relay": (730, 1095), "recloser": (365, 365), "rtu": (365, 365), "batt_q": (30, 91), "batt_a": (365, 365),
           "charger": (365, 365), "cb": (365, 1095), "ct": (1460, 1460), "vt": (1095, 1095)}
FIRMWARE = (("AT-7", "1.9.4", dt.date(2021, 5, 3), dt.date(2021, 12, 10)), ("NX-9", "R3.12", dt.date(2021, 9, 13), dt.date(2022, 6, 30)),
            ("NX-5", "4.2.1", dt.date(2022, 4, 4), dt.date(2023, 1, 27)), ("KDR-30", "2.7.0", dt.date(2023, 2, 6), dt.date(2023, 11, 24)),
            ("NX-9", "R3.16", dt.date(2024, 3, 4), dt.date(2025, 2, 28)), ("RT-600", "6.40", dt.date(2025, 4, 7), dt.date(2025, 12, 12)))
TOKEN = "@@%s-%d-%d@@"


def _crew(rng, region_pool):
    names = [c for c, _ in region_pool]
    w = np.array([p for _, p in region_pool])
    return names[int(rng.choice(len(names), p=w / w.sum()))]


def _hours(rng, crew, mu, sd, lo=0.25, travel=True):
    """task hours at the crew's pace plus the crew's travel to site, booked to the quarter hour"""
    trip, pace = CREW[crew]
    h = float(np.exp(rng.normal(np.log(mu), sd))) * pace
    if travel:
        h += trip * float(np.exp(rng.normal(0.0, 0.3)))
    return max(lo, round(h * 4) / 4)


def _lag(rng, kind):
    """days from raising to completion: calendar days for corrective work, working days for planned work"""
    u = rng.random()
    if kind == "CM":
        return 0 if u < 0.45 else 1 if u < 0.66 else int(rng.integers(2, 5)) if u < 0.82 else int(rng.integers(5, 15)) if u < 0.94 else int(rng.integers(15, 46))
    if kind == "EM":
        return 0 if u < 0.8 else 1 if u < 0.95 else int(rng.integers(2, 4))
    if kind == "INSP":
        return 0 if u < 0.85 else 1 if u < 0.95 else int(rng.integers(2, 6))
    if kind == "MOD":
        return 0 if u < 0.55 else int(rng.integers(1, 6)) if u < 0.85 else int(rng.integers(6, 21)) if u < 0.97 else int(rng.integers(21, 61))
    return 0 if u < 0.6 else int(rng.integers(1, 4)) if u < 0.95 else int(rng.integers(4, 9))


def _slip(rng):
    """working days a planned test runs past the date it falls due: outages, crew availability and storms"""
    u = rng.random()
    return (int(rng.integers(0, 9)) if u < 0.52 else int(rng.integers(9, 22)) if u < 0.76 else
            int(rng.integers(22, 46)) if u < 0.90 else int(rng.integers(46, 111)) if u < 0.97 else int(rng.integers(111, 221)))


def work_orders(subs, reg, f, sp, ships, rebuilt, fits, eq):
    cur = dict(zip(subs["sub_code"], subs["region"]))
    leg = dict(zip(subs["sub_code"], subs["legacy_region"]))
    kind_of = dict(zip(subs["sub_code"], subs["station_type"]))
    aid2sub = dict(zip(reg["asset_id"], reg["sub_code"]))
    aid2sub.update(zip(eq["asset_id"], eq["sub_code"]))
    model = dict(zip(reg["asset_id"], reg["model"]))
    bay_of = dict(zip(reg["asset_id"], reg["bay"]))
    maker = dict(zip(reg["asset_id"], reg["manufacturer"]))
    inst = dict(zip(reg["asset_id"], reg["install_date"]))
    retd = {a: (r if isinstance(r, dt.date) else None) for a, r in zip(reg["asset_id"], reg["retired_date"])}
    key2aid = dict(zip(sp["relay_key"], sp["asset_id"]))
    aid2key = {a: k for k, a in key2aid.items()}
    spell = {(int(k), int(m)): (s, rp) for k, m, s, rp in zip(sp["relay_key"], sp["module"], sp["serial"], sp["repaired"])}
    T = crewtext.Writer(W.SEED * 7 + 3)
    rng_c = np.random.default_rng([W.SEED, 30, 0])     # crews and hours
    rows = []

    def in_service(aid, d):
        r = retd.get(aid)
        return inst[aid] <= d and (r is None or d < r)

    def crew(sub, d):
        return _crew(rng_c, CREWS_CURRENT[cur[sub]] if d >= W.REORG else CREWS_LEGACY[leg[sub]])

    def add(d, aid, sub, wt, comp, desc, done, hours, key=None, cr=None, status=None, due=None, after=0):
        """hours: (mu, sd, lo[, travel]) of the task, or a callable of the crew; done: completion date, None while open"""
        cr = cr or crew(sub, d)
        h = hours(cr) if callable(hours) else _hours(rng_c, cr, *hours)
        rows.append(dict(reported_date=d, asset_id=aid, sub_code=sub, work_type=wt, status=status, failed_component=comp,
                         description=desc(cr) if callable(desc) else desc, labour_hours=h, completed_date=done, crew=cr,
                         region=cur[sub] if d >= W.REORG else leg[sub], _key=key, _after=after))

    storms = {}
    for day, regions in W.STORMS:
        for g in regions:
            storms.setdefault(g, []).append(day)

    def after_storm(sub, d):
        hit = [s for s in storms.get(cur[sub], []) if 0 <= (d - s).days <= 2]
        return hit[0] if hit else None

    def day_name(d):
        return "%s %d" % (d.strftime("%b"), d.day)

    # A. PSU failures on the VPR-340 fleet as recorded from the go-live; the paperwork now and then closes a day or
    # two after the new module went in, and after some call-outs the spare goes in on a follow-up work order
    rng_a = np.random.default_rng([W.SEED, 30, 1])
    fw = f[f["month"] >= GOLIVE].sort_values(["reported", "relay_key"], kind="mergesort")
    for r in fw.itertuples():
        k, m = int(r.relay_key), int(r.module)
        aid = key2aid[k]
        sub = aid2sub[aid]
        store = STORES[cur[sub]]
        nxt_serial, nxt_rep = spell.get((k, m + 1), (None, False))
        fit = fits.get((k, m))
        storm = after_storm(sub, r.reported)
        wt = "EM" if r.em else "CM"
        cr = crew(sub, r.reported)
        text = T.psu_failure(cr, bay_of[aid], store, r.em, storm=day_name(storm) if storm else None,
                             fitted="later" if fit else "same", repaired=nxt_rep, serial=nxt_serial,
                             wo_fit=(TOKEN % ("FIT", k, m)) if (fit and rng_a.random() < 0.5) else None)
        done = r.done
        if rng_a.random() < 0.12:
            late = done + dt.timedelta(days=int(rng_a.integers(1, 3)))
            done = late if late <= W.ASOF else done
        add(r.reported, aid, sub, wt, "PSU", text, done, (1.6, 0.35, 0.75) if fit else (2.2, 0.35, 1.0), key=("F", k, m), cr=cr)
        if fit:
            raised, fitted = fit
            cr2 = crew(sub, raised)
            add(raised, aid, sub, "CM", "", T.psu_fitting(cr2, bay_of[aid], store, nxt_serial, nxt_rep, TOKEN % ("F", k, m)),
                fitted, (1.5, 0.35, 0.5), key=("FIT", k, m), cr=cr2, after=1)
        if rng_a.random() < 0.045:
            add(r.reported, aid, sub, wt, "PSU", T.psu_duplicate(cr, TOKEN % ("F", k, m)), None, lambda c: 0.0, status="CANC",
                after=1)

    # B. failed PSUs shipped to Veldtek under RMA, one logistics work order per box
    lkey = {}
    for g, grp in ships[ships["kind"] == "EAM"].groupby("group", sort=False):
        grp = grp.sort_values(["removed", "relay_key"], kind="mergesort")
        k0, m0 = int(grp["relay_key"].iloc[0]), int(grp["module"].iloc[0])
        aid = key2aid[k0]
        sub = aid2sub[aid]
        ship = grp["ship"].iloc[0]
        cr = crew(sub, ship)
        add(ship, aid, sub, "LOG", "PSU", T.log_psu(cr, list(grp["serial"])), S.add_bdays(ship, _lag(rng_c, "LOG")),
            (0.4 + 0.15 * (len(grp) - 1), 0.4, 0.25, False), key=("L", k0, m0), cr=cr)
        for k, m in zip(grp["relay_key"], grp["module"]):
            lkey[(int(k), int(m))] = ("L", k0, m0)
    months = S.months_between(GOLIVE, END)

    # C. corrective work on every protection device (other than the VPR-340 PSU failures above): monthly at a rate
    # that rises with the device's age; the crew records the failed card, or nothing when no fault is found
    rng_d = np.random.default_rng([W.SEED, 30, 2])
    ids = reg["asset_id"].to_numpy()
    mdl = reg["model"].to_numpy()
    base = np.array([CM_BASE[m] for m in mdl])
    ins = np.array([d.toordinal() for d in reg["install_date"]])
    ret = np.array([d.toordinal() if isinstance(d, dt.date) else 10 ** 7 for d in reg["retired_date"]])
    for mo in months:
        mid = (mo + dt.timedelta(days=14)).toordinal()
        live = (ins <= mid) & (ret > mid)
        age = np.maximum(0.0, (mid - ins) / 365.25)
        k = rng_d.poisson(np.where(live, base * (0.6 + 0.04 * age) / 12.0, 0.0))
        last = (S.month_add(mo, 1) - dt.timedelta(days=1)).day
        for j in np.nonzero(k)[0]:
            for _ in range(int(k[j])):
                aid = ids[j]
                d = dt.date(mo.year, mo.month, int(rng_d.integers(1, last + 1)))
                if d > W.ASOF or not in_service(aid, d):
                    continue
                sub = aid2sub[aid]
                kd = KIND[mdl[j]]
                bay = bay_of[aid]
                a_age = (d - inst[aid]).days / 365.25
                wt = "CM" if rng_d.random() < 0.82 else "EM"
                if rng_d.random() < 0.1:
                    add(d, aid, sub, wt, "", lambda c, kd=kd, bay=bay: T.nff(c, kd, bay),
                        d + dt.timedelta(days=_lag(rng_d, wt)), (1.5, 0.4, 0.5))
                    continue
                comps, p = COMP[kd]
                if mdl[j] == "VPR-340":
                    comps, p = comps[:4], np.array(p[:4]) / sum(p[:4])
                comp = comps[int(rng_d.choice(len(comps), p=p))]
                warranty = a_age < 2.0
                recovered = (not warranty) and a_age >= 25 and mdl[j] in ("VPR-220", "VPR-340") and comp != "PSU"
                wait = (dt.date(2026, 3, 1) <= d <= dt.date(2026, 8, 31) and mdl[j] != "VPR-340"
                        and comp in ("CPU", "IO", "COMMS") and rng_d.random() < 0.15)
                if wait:
                    add(d, aid, sub, wt, comp, lambda c, kd=kd, comp=comp, aid=aid, bay=bay:
                        T.awaiting(c, kd, comp, PART[comp], MAKER[maker[aid]], HOLD[kd], bay), None, (2.0, 0.4, 0.5))
                    continue
                done = d + dt.timedelta(days=_lag(rng_d, wt))
                add(d, aid, sub, wt, comp, lambda c, kd=kd, comp=comp, bay=bay, w=warranty, rc=recovered:
                    T.fault(c, kd, comp, bay, warranty=w, recovered=rc), done, (2.4, 0.45, 0.5))
                ships_back = warranty or (comp in ("CPU", "IO", "COMMS") and mdl[j] in ("NX-9", "NX-5", "KDR-30", "RT-600", "AT-7", "VPR-340")
                                          and rng_d.random() < 0.3)
                if ships_back:
                    d2 = S.add_bdays(done, int(rng_d.integers(2, 11)))
                    if d2 <= W.ASOF:
                        add(d2, aid, sub, "LOG", comp, lambda c, comp=comp, aid=aid, w=warranty:
                            T.log_part(c, PART[comp], MAKER[maker[aid]], w), S.add_bdays(d2, _lag(rng_d, "LOG")), (0.4, 0.4, 0.25, False))

    # D. corrective work on the station DC systems, breakers and VTs
    rng_e = np.random.default_rng([W.SEED, 30, 3])
    eqr = {a: (r if isinstance(r, dt.date) else None) for a, r in zip(eq["asset_id"], eq["retired"])}
    for e in eq.itertuples():
        start = max(e.install, GOLIVE)
        stop = min(e.retired or END, END)
        yrs = max(0.0, (stop - start).days / 365.25)
        for _ in range(int(rng_e.poisson(EQ_CM[e.cls] * yrs))):
            d = start + dt.timedelta(days=int(rng_e.integers(0, max(1, (stop - start).days))))
            if d > W.ASOF:
                continue
            wt = "CM" if rng_e.random() < 0.85 else "EM"
            desc = ((lambda c, e=e: T.dc_fault(c, e.cls, e.name)) if e.cls in ("battery", "charger") else
                    (lambda c, e=e: T.cb_fault(c, e.name)) if e.cls == "cb" else (lambda c, e=e: T.vt_fault(c, e.name)))
            add(d, e.asset_id, e.sub_code, wt, "", desc, d + dt.timedelta(days=_lag(rng_e, wt)), (2.0, 0.45, 0.5))

    # E. periodic tests under P&C Standard 3.2. Each test work order is generated on the date the test falls due and
    # closed when the test is done; the next one falls due a full interval after that
    rng_p = np.random.default_rng([W.SEED, 30, 4])

    def schedule(aid, sub, cls, install, retired, mu, note_p=0.35, **kw):
        ts = kind_of[sub] == "TS"
        days = PM_DAYS[cls][0 if ts else 1]
        due = max(GOLIVE + dt.timedelta(days=int(rng_p.integers(0, days))), install + dt.timedelta(days=days))
        end = retired or END
        while True:
            due = S.add_bdays(due, 0)
            if due >= end or due > W.ASOF:
                return
            done = S.add_bdays(due, _slip(rng_p) if cls not in ("batt_q",) else min(_slip(rng_p), 20))
            if retired is not None and done >= retired:
                add(due, aid, sub, "PM", "", "Device retired, test not required", None, lambda c: 0.0, status="CANC")
                return
            add(due, aid, sub, "PM", "", lambda c, cls=cls, due=due: T.pm(c, cls, due, note_p=note_p, **kw), done,
                (mu, 0.3, 0.5))
            if done > W.ASOF:
                return
            due = done + dt.timedelta(days=days)

    for a in reg.itertuples():
        kd = KIND[a.model]
        cls = "rtu" if kd == "rtu" else "recloser" if kd == "recloser" else "relay"
        mu = 4.5 if a.model in ("NX-5", "KDR-30") else 5.0 if kd == "rtu" else 3.0 if kd == "recloser" else 3.5
        schedule(a.asset_id, a.sub_code, cls, a.install_date,
                 a.retired_date if isinstance(a.retired_date, dt.date) else None, mu, bay=a.bay, model=a.model)
    for e in eq.itertuples():
        if e.cls == "battery":
            schedule(e.asset_id, e.sub_code, "batt_q", e.install, e.retired, 1.0, note_p=0.3, name=e.name)
            schedule(e.asset_id, e.sub_code, "batt_a", e.install, e.retired, 3.0, note_p=0.4, name=e.name)
        elif e.cls == "charger":
            schedule(e.asset_id, e.sub_code, "charger", e.install, e.retired, 1.2, note_p=0.3, name=e.name)
        elif e.cls == "cb":
            schedule(e.asset_id, e.sub_code, "cb", e.install, e.retired, 1.5, note_p=0.35, bay=e.name)
            schedule(e.asset_id, e.sub_code, "ct", e.install, e.retired, 2.5, note_p=0.3, bay=e.name)
        else:
            schedule(e.asset_id, e.sub_code, "vt", e.install, e.retired, 1.2, note_p=0.3, name=e.name)

    # F. station inspections by the P&C crews: monthly at transformer stations and quarterly at distribution stations
    # (a few missed), relay event records reviewed twice a year and an infrared scan of the panels yearly at the
    # transformer stations, and checks after the storms that brought crews out across a region
    rng_i = np.random.default_rng([W.SEED, 30, 5])
    for s in subs.itertuples():
        ts = s.station_type == "TS"
        phase = int(rng_i.integers(0, 3))
        ir_month = int(rng_i.integers(3, 11))
        for mo in months:
            if not ts and (mo.month - 1) % 3 != phase:
                continue
            if rng_i.random() < 0.03:
                continue
            d = S.pick_bday(rng_i, mo, mo + dt.timedelta(days=24))
            if d > W.ASOF:
                continue
            add(d, "", s.sub_code, "INSP", "", lambda c, d=d, ts=ts: T.insp(c, "ts" if ts else "ds", d),
                S.add_bdays(d, _lag(rng_i, "INSP")), (0.6 if ts else 0.7, 0.35, 0.5))
            if ts and mo.month in (3, 9) and rng_i.random() < 0.9:
                d5 = S.pick_bday(rng_i, d + dt.timedelta(days=1), d + dt.timedelta(days=10))
                if d5 <= W.ASOF:
                    add(d5, "", s.sub_code, "INSP", "", lambda c, d5=d5: T.insp(c, "events", d5), S.add_bdays(d5, _lag(rng_i, "INSP")),
                        (1.5, 0.3, 0.5))
            if ts and mo.month == ir_month:
                d3 = S.pick_bday(rng_i, d + dt.timedelta(days=1), d + dt.timedelta(days=10))
                if d3 <= W.ASOF:
                    add(d3, "", s.sub_code, "INSP", "", lambda c, d3=d3: T.insp(c, "ir", d3), S.add_bdays(d3, _lag(rng_i, "INSP")),
                        (1.5, 0.3, 0.5))
    # settings audits at the transformer stations: each relay's as-left settings against the approved file every
    # two years
    for a in reg[(reg["sub_code"].map(kind_of) == "TS") & reg["model"].map(lambda m: KIND[m] == "relay")].itertuples():
        end = a.retired_date if isinstance(a.retired_date, dt.date) else END
        d = max(GOLIVE + dt.timedelta(days=int(rng_i.integers(0, 730))), a.install_date + dt.timedelta(days=365))
        while True:
            d = S.add_bdays(d, 0)
            if d >= end or d > W.ASOF:
                break
            add(d, a.asset_id, a.sub_code, "INSP", "", lambda c, d=d, b=a.bay, m=a.model: T.audit(c, b, m, d),
                S.add_bdays(d, _lag(rng_i, "INSP")), (1.2, 0.3, 0.5))
            d = d + dt.timedelta(days=730 + int(rng_i.integers(-20, 41)))
    for day, regions in W.STORMS:
        for s in subs.itertuples():
            if cur[s.sub_code] in regions and rng_i.random() < 0.35:
                d = day + dt.timedelta(days=int(rng_i.integers(0, 2)))
                add(d, "", s.sub_code, "INSP", "", lambda c, day=day: T.storm_check(c, day_name(day)),
                    d + dt.timedelta(days=int(rng_i.random() < 0.15)), (1.0, 0.35, 0.5))

    # G. modifications: settings changes from protection coordination studies (a station at a time), vendor firmware
    # campaigns, settings reissued after feeder reconfiguration, a few held open
    rng_m = np.random.default_rng([W.SEED, 30, 6])
    by_sub = {sc: list(g["asset_id"]) for sc, g in reg[reg["model"].map(lambda m: KIND[m] == "relay")].groupby("sub_code")}
    subs_list = sorted(by_sub)
    weights = np.array([len(by_sub[sc]) for sc in subs_list], dtype=float)
    seq = {}
    for mo in months:
        for _ in range(int(rng_m.poisson(7))):
            sc = subs_list[int(rng_m.choice(len(subs_list), p=weights / weights.sum()))]
            yy = mo.year % 100
            seq[yy] = seq.get(yy, 0) + int(rng_m.integers(1, 4))
            study = "PCS-%02d-%03d" % (yy, seq[yy])
            live = [a for a in by_sub[sc] if in_service(a, mo)]
            if not live:
                continue
            pick = rng_m.choice(len(live), min(len(live), int(rng_m.integers(1, 9))), replace=False)
            start = S.pick_bday(rng_m, mo, S.month_add(mo, 1) - dt.timedelta(days=1))
            for j in sorted(pick):
                aid = live[int(j)]
                d = S.add_bdays(start, int(rng_m.integers(0, 15)))
                if d > W.ASOF or not in_service(aid, d):
                    continue
                if dt.date(2026, 4, 1) <= d <= dt.date(2026, 8, 31) and rng_m.random() < 0.02:
                    add(d, aid, sc, "MOD", "", lambda c, st=study: T.held(c, st), None, (1.6, 0.4, 0.25))
                    continue
                add(d, aid, sc, "MOD", "", lambda c, b=bay_of[aid], st=study, d=d: T.settings(c, b, st, d),
                    S.add_bdays(d, _lag(rng_m, "MOD")), (1.6, 0.4, 0.25))
        for _ in range(int(rng_m.poisson(3))):
            aid = ids[int(rng_m.integers(len(ids)))]
            d = S.pick_bday(rng_m, mo, S.month_add(mo, 1) - dt.timedelta(days=1))
            if d > W.ASOF or not in_service(aid, d) or KIND[model[aid]] != "relay":
                continue
            add(d, aid, aid2sub[aid], "MOD", "", lambda c, b=bay_of[aid]: T.reissue(c, b), S.add_bdays(d, _lag(rng_m, "MOD")),
                (1.4, 0.4, 0.25))
    for mdl_name, rel, a0, a1 in FIRMWARE:
        span = S.bdays(a0, a1)
        for a in reg[reg["model"] == mdl_name].itertuples():
            if rng_m.random() > 0.82:
                continue
            d = span[int(rng_m.integers(len(span)))]
            if d > W.ASOF or not in_service(a.asset_id, d):
                continue
            add(d, a.asset_id, a.sub_code, "MOD", "", lambda c, rel=rel: T.firmware(c, rel), S.add_bdays(d, _lag(rng_m, "MOD")),
                (1.2, 0.4, 0.25))

    # replacements: every retirement is a cutover to the new device in the same bay on the same day
    newdev = {(s, b, d): mo for s, b, d, mo in zip(reg["sub_code"], reg["bay"], reg["install_date"], reg["model"])}
    for a in reg[reg["status"] == "RETIRED"].itertuples():
        d1 = a.retired_date
        nm = newdev.get((a.sub_code, a.bay, d1), a.model)
        if a.sub_code in rebuilt:
            t = "Rebuild cutover %s: %s removed, %s commissioned" % (a.bay, a.model, nm)
            if a.bay == "RTU":
                t += ", SCADA points verified"
        elif a.bay == "RTU":
            t = "RTU replacement: %s removed, %s commissioned, point-to-point test with control room" % (a.model, nm)
        else:
            t = "Replace %s on %s with %s, commissioned and trip tested" % (a.model, a.bay, nm)
        d0 = S.pick_bday(rng_m, max(GOLIVE, d1 - dt.timedelta(days=21)), d1)
        key = ("R", int(aid2key[a.asset_id])) if a.asset_id in aid2key else None
        add(d0, a.asset_id, a.sub_code, "MOD", "", t, d1, (7.0 if a.bay == "RTU" else 4.5, 0.3, 2.0), key=key)
    for e in eq[[isinstance(r, dt.date) for r in eq["retired"]]].itertuples():
        d1 = e.retired
        t = ("Replace battery bank %s, new VRLA bank installed, capacity test OK" % e.name if e.cls == "battery" else
             "Replace charger %s, new unit commissioned, alarms tested" % e.name)
        d0 = S.pick_bday(rng_m, max(GOLIVE, d1 - dt.timedelta(days=21)), d1)
        add(d0, e.asset_id, e.sub_code, "MOD", "", t, d1, (6.0, 0.3, 2.0))

    # work raised and then cancelled across every type (raised in error, superseded, access refused)
    rng_x = np.random.default_rng([W.SEED, 30, 7])
    plain = [k for k, r in enumerate(rows) if r["_key"] is None and r["status"] is None]
    for k in rng_x.choice(plain, int(len(rows) * 0.012), replace=False):
        r = dict(rows[int(k)])
        r.update(status="CANC", completed_date=None, labour_hours=0.0, _key=None,
                 description=["Raised in error", "Superseded by planned outage work", "No access - rebooked",
                              "Duplicate of earlier work order", "Raised against wrong asset", "Combined with outage work",
                              "No outage available, rebooked"][int(rng_x.integers(7))])
        rows.append(r)

    # status at the extract: open work is in progress; completed work is closed after supervisor review, which takes
    # days to weeks, and a little older work is never closed
    for r in rows:
        if r["status"] == "CANC":
            r["completed_date"] = None
            continue
        done = r["completed_date"]
        if done is None or done > W.ASOF:
            r["status"], r["completed_date"] = "INPRG", None
        elif rng_x.random() < 0.012:
            r["status"] = "COMP"
        else:
            review = 1 + float(rng_x.gamma(1.6, 7.0))
            r["status"] = "CLOSE" if done + dt.timedelta(days=review) <= W.ASOF else "COMP"
    wo = pd.DataFrame(rows)
    wo = wo[wo["reported_date"] <= W.ASOF]
    wo = wo.sort_values(["reported_date", "sub_code", "asset_id", "_after", "work_type", "status", "description"],
                        kind="mergesort").reset_index(drop=True)
    # EAM work-order numbers run in raise order and are shared with the lines and metering crews
    n, nums = 2104000, []
    for _ in range(len(wo)):
        n += int(rng_x.geometric(0.28)) + (int(rng_x.integers(20, 300)) if rng_x.random() < 0.004 else 0)
        nums.append("WO%07d" % n)
    wo.insert(0, "wo_number", nums)
    keymap = {k: w for k, w in zip(wo["_key"], wo["wo_number"]) if k is not None and not (isinstance(k, float))}
    for k, lk in lkey.items():
        keymap[("L",) + k] = keymap[lk]

    def detok(t):
        while "@@" in t:
            i = t.index("@@")
            j = t.index("@@", i + 2)
            kind, a, b = t[i + 2:j].split("-")
            t = t[:i] + keymap[(kind, int(a), int(b))] + t[j + 2:]
        return t

    wo["description"] = [detok(t) if "@@" in t else t for t in wo["description"]]
    cols = ["wo_number", "reported_date", "completed_date", "work_type", "status", "asset_id", "sub_code", "region",
            "failed_component", "description", "crew", "labour_hours"]
    return wo[cols], keymap


def module_history(sp, keymap, fits, yearonly):
    """EAM rotating-item installation history for VPR340-PSU-02: every module in a VPR-340 position since the
    go-live (modules in place at go-live were loaded with the fitted date the legacy system held, which for the
    oldest relays was sometimes only the year, loaded as 1 July)"""
    h = sp[(sp["removed"].isna()) | (sp["removed"] >= GOLIVE)].copy()
    h["removal_wo"] = [keymap.get(("F", int(k), int(m)), "") if r == "FAILED" else keymap.get(("R", int(k)), "") if r == "RETIRED"
                       else "" for k, m, r in zip(h["relay_key"], h["module"], h["reason"])]
    fit_wo, fit_date = [], []
    for k, m, f in zip(h["relay_key"], h["module"], h["fitted"]):
        k, m = int(k), int(m)
        if m > 0 and f >= GOLIVE:
            fp = fits.get((k, m - 1))
            fit_wo.append(keymap[("FIT", k, m - 1)] if fp else keymap.get(("F", k, m - 1), ""))
            fit_date.append(fp[1] if fp else f)
        else:
            fit_wo.append("")
            fit_date.append(dt.date(f.year, 7, 1) if (m == 0 and k in yearonly and f < YEAR_ONLY_BEFORE) else f)
    h["fitting_wo"] = fit_wo
    h["fitted_shown"] = fit_date
    h["removal_reason"] = [{"FAILED": "FAILED", "RETIRED": "ASSET RETIRED"}.get(r, "") if isinstance(r, str) else "" for r in h["reason"]]
    h["record_source"] = ["GO-LIVE LOAD" if f < GOLIVE else "EAM" for f in h["fitted"]]
    h["item_code"] = "VPR340-PSU-02"
    out = pd.DataFrame(dict(serial_no=h["serial"], item_code=h["item_code"], asset_id=h["asset_id"],
                            fitted_date=h["fitted_shown"], fitting_wo=h["fitting_wo"],
                            removed_date=[None if (r is None or pd.isna(r)) else r for r in h["removed"]],
                            removal_reason=h["removal_reason"], removal_wo=h["removal_wo"],
                            record_source=h["record_source"]))
    return out.sort_values(["asset_id", "fitted_date"], kind="mergesort").reset_index(drop=True)


def stock(n_open):
    """spares on hand at the count, all P&C items; the VPR-340 PSU lines are the ones the buy uses"""
    psu = ("300-14522", "VPR340-PSU-02", "Power supply module, VPR-340, 48-250 V DC")
    rows = [psu + (loc, b, "SERVICEABLE", "EA", q, 1418.40) for loc, b, q in W.SERVICEABLE]
    rows += [psu + (loc, b, "QUARANTINE", "EA", q, 1418.40) for loc, b, q in W.QUARANTINE]
    rows += [psu + ("Veldtek repair centre", "", "AT REPAIR", "EA", n_open, 1418.40)]
    rows += [
        ("300-14523", "VPR340-CPU-01", "Processor card, VPR-340", "Central Stores (Barrie Road)", "A-14-04", "SERVICEABLE", "EA", 6, 2210.00),
        ("300-14524", "VPR340-IO-03", "Input/output card, VPR-340", "Central Stores (Barrie Road)", "A-14-05", "SERVICEABLE", "EA", 9, 865.75),
        ("300-14526", "VPR340-DSP-01", "Front panel display, VPR-340", "North Depot", "N-02-12", "SERVICEABLE", "EA", 4, 640.00),
        ("300-11870", "VPR220-PSU-01", "Power supply module, VPR-220", "Central Stores (Barrie Road)", "A-11-01", "SERVICEABLE", "EA", 12, 905.20),
        ("300-11870", "VPR220-PSU-01", "Power supply module, VPR-220", "Central Stores (Barrie Road)", "Q-01-04", "QUARANTINE", "EA", 2, 905.20),
        ("300-20115", "NX5-PSU", "Power supply, Norland NX-5", "Central Stores (Barrie Road)", "A-20-02", "SERVICEABLE", "EA", 7, 1290.00),
        ("300-20131", "NX9-PSU", "Power supply, Norland NX-9", "North Depot", "N-04-03", "SERVICEABLE", "EA", 5, 1105.00),
        ("300-31007", "AT7-CTRL", "Control module, Arcwell AT-7", "South Depot", "S-01-09", "SERVICEABLE", "EA", 3, 3480.00),
        ("300-26402", "KDR30-PSU", "Power supply, Kestrel KDR-30", "North Depot", "N-03-02", "SERVICEABLE", "EA", 2, 1675.50),
        ("300-31040", "RT400-CPU", "RTU processor, Arcwell RT-400", "Central Stores (Barrie Road)", "A-31-01", "SERVICEABLE", "EA", 2, 4120.00),
    ]
    return pd.DataFrame(rows, columns=["item_no", "item_code", "item_description", "location", "bin", "stock_state", "uom",
                                       "quantity", "unit_cost_cad"])


SERVICE_NOTES = {
    "REPAIRED": ["Electrolytic caps and fan replaced, load tested", "Caps C12-C19 and fan replaced, 4 h burn-in OK",
                 "Bulk caps bulging, replaced with fan", "Fan seized, caps replaced, tested OK", "Replaced caps, fan, output fuse",
                 "Caps and fan replaced", "Output stage repaired, caps and fan replaced", "", ""],
    "NO FAULT FOUND": ["NFF on bench, 24 h burn-in OK", "No fault found, returned", "Tested at full load, no fault"],
    "BEYOND ECONOMIC REPAIR": ["Board burnt, BER, credit issued", "Transformer failed, not economic, scrapped",
                               "PCB delaminated, BER", "BER, scrapped at customer's request"],
    "OPEN": ["In repair", "Awaiting caps", "On bench", "", ""],
    "SCRAPPED (SB-15-02)": ["Returned under SB-15-02, scrapped", "SB-15-02 exchange, credit issued", "SB-15-02 return"],
}


def rma_log(ships, keymap):
    """Veldtek's RMA register for VPR340-PSU-02 returns from Tarrock. Veldtek's numbering runs across customers, so
    the numbers climb with the date each RMA was requested; Tarrock's reference is whatever the stores clerk quoted."""
    rr = np.random.default_rng([W.SEED, 22])
    s = ships.copy()
    s["req"] = [S.add_bdays(d - dt.timedelta(days=int(rr.integers(1, 5))), 0) for d in s["ship"]]
    s = s.sort_values(["req", "group", "removed", "relay_key"], kind="mergesort").reset_index(drop=True)
    nums, last = [], 0
    for d in s["req"]:
        base = 11800 + int((d - dt.date(2015, 1, 1)).days * 1.85) + int(rr.integers(0, 6))
        last = max(last + 1, base)
        nums.append("RMA-%05d" % last)
    refs, cust = [], []
    for r in s.itertuples():
        if r.kind == "EAM":
            wo = keymap[("L", int(r.relay_key), int(r.module))]
            u = rr.random()
            if u < 0.06:
                wo = ""
            elif u < 0.075:
                wo = wo[2:]
            elif u < 0.09:
                i = int(rr.integers(3, len(wo) - 1))
                wo = wo[:i] + wo[i + 1] + wo[i] + wo[i + 2:]
            elif u < 0.10:
                wo = "WO " + wo[2:]
            elif u < 0.11:
                wo = keymap[("F", int(r.relay_key), int(r.module))]
            refs.append(wo)
        else:
            refs.append("" if rr.random() < 0.08 else ("SB-15-02 / " + r.ref if r.kind == "SB" and rr.random() < 0.5 else r.ref))
        y = r.received.year
        cust.append("Tarrock Hydro" if y < 2019 else ("TARROCK HYDRO DIST" if rr.random() < 0.04 else "Tarrock Hydro Distribution"))
    notes = [SERVICE_NOTES[res][int(rr.integers(len(SERVICE_NOTES[res])))] for res in s["result"]]
    return pd.DataFrame(dict(rma_no=nums, part_no="VPR340-PSU-02", serial_no=s["serial"], customer=cust, customer_ref=refs,
                             received=s["received"], result=s["result"], returned=s["returned"], notes=notes))


def rma_xlsx(path, rma):
    """the register as Veldtek's service centre keeps it in Excel"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RMA register"
    hdr = ["RMA No.", "Part No.", "Serial No.", "Customer", "Customer Ref.", "Date Received", "Result", "Date Returned",
           "Service notes"]
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
    for r in rma.itertuples(index=False):
        ws.append([r.rma_no, r.part_no, r.serial_no, r.customer, r.customer_ref or None, r.received, r.result,
                   r.returned if isinstance(r.returned, dt.date) else None, r.notes or None])
    for row in ws.iter_rows(min_row=2):
        for j in (5, 7):
            if row[j].value is not None:
                row[j].number_format = "dd-mmm-yyyy"
    for j, w in enumerate([12, 15, 16, 28, 15, 14, 24, 14, 40], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:I%d" % ws.max_row
    wb.save(path)
