"""The pack's documents, in their owners' voices."""
import datetime as dt
import os
import shutil
import subprocess
import tempfile

import docx
import numpy as np
import openpyxl
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl.styles import Alignment, Font, PatternFill

import sim as S
import world as W


def _doc(font="Calibri", size=10.5, author="", title="", when=None):
    d = docx.Document()
    cp = d.core_properties
    cp.author = cp.last_modified_by = author
    cp.title = title
    cp.comments = cp.keywords = cp.subject = cp.category = ""
    if when is not None:
        cp.created = cp.modified = when
    cp.revision = 1
    lang = d.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr")).find(qn("w:lang"))
    if lang is not None:
        lang.set(qn("w:val"), "en-CA")
    st = d.styles["Normal"]
    st.font.name = font
    st.font.size = Pt(size)
    for name, sz in (("Heading 1", 13), ("Heading 2", 11)):
        h = d.styles[name]
        h.font.name = font
        h.font.size = Pt(sz)
        h.font.bold = True
        h.font.color.rgb = RGBColor(0x1F, 0x2A, 0x33)
        h.paragraph_format.space_before = Pt(10 if name == "Heading 2" else 4)
        h.paragraph_format.space_after = Pt(3)
        rpr = h.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(a), font)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if fonts.get(qn(a)) is not None:
                del fonts.attrib[qn(a)]
    for s in d.sections:
        s.page_width, s.page_height = Cm(21.59), Cm(27.94)
        s.left_margin = s.right_margin = Cm(2.3)
        s.top_margin = Cm(1.9)
        s.bottom_margin = Cm(1.8)
    return d


def _p(d, text, bold=False, size=None, after=4, italic=False, style=None, align=None):
    p = d.add_paragraph(style=style)
    r = p.add_run(text)
    if bold:
        r.bold = True
    if italic:
        r.italic = True
    if size:
        r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(after)
    if align:
        p.alignment = align
    return p


def _field(par, instr, size):
    """a simple field the way Word stores it: begin, instruction, separate, cached result and end each in a run"""
    for kind in ("begin", "instr", "separate", "result", "end"):
        r = par.add_run()
        r.font.size = size
        if kind == "instr":
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = " %s " % instr
        elif kind == "result":
            e = OxmlElement("w:t")
            e.text = "1"
        else:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), kind)
        r._r.append(e)


def _footer(d, text):
    f = d.sections[0].footer.paragraphs[0]
    r = f.add_run(text + "    Page ")
    r.font.size = Pt(8)
    _field(f, "PAGE", Pt(8))
    r2 = f.add_run(" of ")
    r2.font.size = Pt(8)
    _field(f, "NUMPAGES", Pt(8))


def _cell(c, text, bold=False, size=9):
    for r in list(c.paragraphs[0].runs):
        r._r.getparent().remove(r._r)
    c.paragraphs[0].paragraph_format.space_after = Pt(0)
    if not text:
        return
    r = c.paragraphs[0].add_run(text)
    if bold:
        r.bold = True
    r.font.size = Pt(size)


def _widths(t, cms):
    """column widths in the table grid and in every cell, as Word keeps them"""
    t.autofit = False
    for col, w in zip(t.columns, cms):
        col.width = Cm(w)
    for row in t.rows:
        for c, w in zip(row.cells, cms):
            c.width = Cm(w)


def _pdf(d, out_path, stem):
    tmp = tempfile.mkdtemp()
    src = os.path.join(tmp, stem + ".docx")
    d.save(src)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, src],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    shutil.move(os.path.join(tmp, stem + ".pdf"), out_path)
    shutil.rmtree(tmp)


def procedure_docx(path):
    d = _doc(author="Kurt Allison", title="AM-PC-014 Obsolete protection spares: final-buy sizing")
    _p(d, "Tarrock Hydro Distribution  |  Asset Management, Protection & Control", size=8.5, italic=True)
    _p(d, "AM-PC-014  Obsolete protection spares: final-buy sizing", style="Heading 1", after=6)
    t = d.add_table(rows=3, cols=4)
    t.style = "Table Grid"
    for r, row in enumerate([("Document", "AM-PC-014", "Revision", "4"),
                             ("Owner", "P&C Asset Planning", "Issued", "14 April 2026"),
                             ("Approver", "Director, Asset Management", "Next review", "April 2028")]):
        for c, v in enumerate(row):
            _cell(t.rows[r].cells[c], v, bold=c in (0, 2))
    _p(d, "", after=2)
    sections = [
        ("1. Purpose",
         ["When a manufacturer withdraws a replaceable module of a protection relay we place one last-time buy, and "
          "failed modules are replaced from that buy for as long as the relays stay in service. This procedure sets how "
          "the buy is sized so that the relays can be kept running until the Protection Renewal Programme replaces them."]),
        ("2. Scope",
         ["Replaceable modules of protection relays and controllers that are covered by a manufacturer's end-of-life "
          "notice and are tracked as rotating items in EAM. Whole-relay spares are held under AM-PC-009 and are not "
          "sized here."]),
        ("3. Definitions",
         ["Module age: the age of a module in completed years on the first day of a month, counted from the date it was "
          "fitted to the relay as recorded in the EAM rotating-item history.",
          "In service: a module counts as in service for a month if it is fitted to a relay in service on the first day "
          "of that month.",
          "Lot: the manufacturer's minimum order quantity stated in the end-of-life notice."]),
        ("4. Support period",
         ["4.1 A relay is supported from the stock count date until the scheme that replaces it is commissioned. Use the "
          "approved Protection Renewal Programme current at the stock count date. Where a scheme is built in stages, the "
          "relay's bay decides which stage replaces it. A relay is counted for every month up to, but not including, the "
          "commissioning month of the scheme or stage that replaces its bay.",
          "4.2 A relay at a substation with no scheme in the programme period is supported to the end of the programme "
          "period (December 2035 in the 2026 programme)."]),
        ("5. Failure rates",
         ["5.1 Module failures are corrective work orders (work type CM or EM) on a relay of the type concerned whose "
          "failed component is the module, with status COMP or CLOSE. Cancelled work orders and logistics work orders "
          "(work type LOG) are not failures.",
          "5.2 Failure rates are set by module age band: 0 to 9, 10 to 14, 15 to 19, 20 to 24, and 25 years and over."
          " A failure takes the age of the failed module on the first day of the month in which the work order was "
          "reported.",
          "5.3 The failure rate for a band is the number of module failures in the work-order history in the band divided "
          "by the module-years in service in the band, counted in module-months."]),
        ("6. Expected failures",
         ["6.1 A module in service for a month is expected to fail its band rate divided by twelve times in that month.",
          "6.2 Expected failures over the support period are the band rates applied to the module-months expected in each "
          "band, counted for every supported relay from the stock count date to the end of its support."]),
        ("7. Available modules",
         ["Available modules are those serviceable on hand at all stores and depots, plus modules at the manufacturer for "
          "repair that the end-of-life notice says will be returned. Quarantined modules are not available until they pass "
          "bench test and are left out of the count."]),
        ("8. The buy",
         ["8.1 Module demand over the support period is treated as Poisson with mean equal to the expected failures.",
          "8.2 Buy the smallest whole number of lots for which the probability that available modules plus the buy cover "
          "all demand over the support period is at least 95 per cent.",
          "8.3 The sizing goes to the Director, Asset Management for signature with the lot count, the cover it gives, the "
          "expected failures by year and the cost."]),
        ("9. Roles",
         ["9.1 The Manager, P&C Asset Planning owns this procedure, prepares the sizing when an end-of-life notice is "
          "received and submits it for signature.",
          "9.2 The Senior Asset Engineer, P&C runs the EAM extracts the sizing uses and keeps them with the sizing.",
          "9.3 The Stores Supervisor confirms the stock count, including modules in quarantine and modules at the "
          "manufacturer under an open RMA.",
          "9.4 The Director, Asset Management approves the buy and signs the purchase requisition."]),
        ("10. Records",
         ["The signed sizing, the end-of-life notice and the EAM extracts are filed in the Asset Management document "
          "library under the item code, and kept for the life of the relays concerned and seven years after."]),
        ("11. Related documents",
         ["AM-PC-009 Protection relay spares holdings",
          "P&C Standard 3.2 Protection relay maintenance and testing",
          "Protection Renewal Programme, approved version",
          "SC-PR-07 Purchasing: last-time buys and single-source items"]),
    ]
    for head, paras in sections:
        _p(d, head, style="Heading 2", after=3)
        for tx in paras:
            _p(d, tx, after=5 if not head.startswith("11.") else 1)
    _p(d, "Approval", style="Heading 2", after=3)
    t = d.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for c, v in enumerate(("", "Name", "Title", "Date")):
        _cell(t.rows[0].cells[c], v, bold=True)
    for row in (("Prepared", "Kurt Allison", "Senior Asset Engineer, P&C", "7 April 2026"),
                ("Reviewed", "Rachael Villanueva", "Manager, P&C Asset Planning", "9 April 2026"),
                ("Approved", "Angel Baird", "Director, Asset Management", "14 April 2026")):
        cells = t.add_row().cells
        for c, v in enumerate(row):
            _cell(cells[c], v, bold=c == 0)
    _p(d, "", after=2)
    _p(d, "Revision history", style="Heading 2", after=3)
    t = d.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for c, v in enumerate(("Rev", "Date", "Change")):
        _cell(t.rows[0].cells[c], v, bold=True)
    for row in (("1", "May 2019", "First issue (VPS-120 display module buy)."),
                ("2", "Mar 2022", "EAM work orders become the failure source after the January 2021 go-live."),
                ("3", "Jun 2023", "Poisson cover test replaces the 20 per cent contingency."),
                ("4", "Apr 2026", "Rates by module age from the EAM rotating-item history; support by bay for the staged "
                                  "schemes of the approved 2026 programme; roles, records and approval added.")):
        cells = t.add_row().cells
        for c, v in enumerate(row):
            _cell(cells[c], v)
    _widths(d.tables[0], (3.0, 5.0, 3.0, 5.0))
    _widths(d.tables[1], (2.4, 4.0, 6.0, 3.6))
    _widths(d.tables[2], (1.2, 2.4, 12.4))
    _footer(d, "AM-PC-014 rev 4    Uncontrolled when printed")
    d.save(path)


def eol_pdf(path):
    d = _doc(font="Arial", size=10, author="George Ray", title="Product end-of-life notice EOL-2026-07")
    _p(d, "VELDTEK PROTECTION SYSTEMS INC.", bold=True, size=11, after=0)
    _p(d, "Lifecycle Services  \u00b7  88 Cardinal Gate, Unit 4, Mississauga ON  L5T 2W8", size=8, after=10)
    _p(d, "Product end-of-life notice EOL-2026-07", style="Heading 1", after=2)
    _p(d, "20 August 2026", after=10)
    _p(d, "To: customers operating VPR-340 feeder protection relays", after=2)
    _p(d, "Re: VPR340-PSU-02 power supply module, last-time buy and end of repair", bold=True, after=8)
    body = [
        "The VPR-340 relay was withdrawn from sale in 2018. Its power supply module, part VPR340-PSU-02, depends on "
        "components that our suppliers have now discontinued, and we will build one final production run.",
        "Last-time-buy orders for VPR340-PSU-02 must reach us by 15 December 2026. Modules are sold in lots of 25 at "
        "CAD 1,965.00 per module (CAD 49,125.00 per lot), delivery by 31 March 2027. No orders are accepted after the "
        "final run.",
        "Repair service for VPR340-PSU-02 ends on 15 December 2026. Modules already at our repair centre under an open "
        "RMA, and modules received under an RMA by that date, will be repaired or replaced and returned. Modules "
        "received after 15 December 2026 will be sent back unrepaired.",
    ]
    for tx in body:
        _p(d, tx, after=6)
    t = d.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    for c, v in enumerate(("Key dates", "")):
        _cell(t.rows[0].cells[c], v, bold=True, size=9.5)
    for k, v in (("Notice issued", "20 August 2026"), ("Last day for last-time-buy orders", "15 December 2026"),
                 ("Last day for RMA receipts at the repair centre", "15 December 2026"),
                 ("Final shipment of last-time-buy orders", "By 31 March 2027")):
        cells = t.add_row().cells
        _cell(cells[0], k, size=9.5)
        _cell(cells[1], v, size=9.5)
    _widths(t, (9.0, 5.0))
    _p(d, "", after=4)
    for tx in [
        "The processor, input/output and display cards of the VPR-340 remain available until further notice. Customers "
        "planning feeder protection replacement can migrate VPR-340 settings to the VPR-560 platform with our conversion "
        "tool; application notes are on the customer portal.",
        "For quotations and order acknowledgements please contact your Veldtek account manager or open a Lifecycle case on "
        "the customer portal, quoting EOL-2026-07.",
    ]:
        _p(d, tx, after=6)
    _p(d, "", after=4)
    _p(d, "George Ray", after=0)
    _p(d, "Product Lifecycle Manager, Protection Relays", after=0)
    _p(d, "Veldtek Protection Systems Inc.", after=0)
    _pdf(d, path, "eol")


SCOPE_ONE = ["Feeder protection replacement %s, new P&C panels", "Replace feeder relays %s and RTU, IEC 61850 station bus",
             "P&C rebuild with new control building, feeders %s", "Feeder protection and control replacement %s",
             "Replace feeder protection %s, panel rewiring and DC upgrade"]
SCOPE_S1 = ["Stage 1: feeder protection %s, new control building", "Stage 1: replace feeder relays %s, station bus",
            "Stage 1: P&C panels %s"]
SCOPE_S2 = ["Stage 2: feeder protection %s, RTU cutover", "Stage 2: replace feeder relays %s", "Stage 2: P&C panels %s, decommission old building"]
SCOPE_VPR220 = ["Replace VPR-220 feeder protection %s", "Replace VPR-220 relays %s, new DC system",
                "VPR-220 feeder relay replacement %s, panel rewiring"]


def other_schemes(rng, prog, reg):
    """Schemes at stations without VPR-340: the feeder bays still on VPR-220 relays and RTUs installed before 2009,
    the oldest equipment first; the stations whose equipment is youngest wait beyond the period"""
    hosts = set(prog["sub_code"])
    live = reg[(reg["status"] == "IN SERVICE") & ~reg["sub_code"].isin(hosts)]
    rows = []
    for sc, g in live.groupby("sub_code"):
        v = g[g["model"] == "VPR-220"]
        rtu = g[(g["model"] == "RT-400") & (g["install_date"] < dt.date(2009, 1, 1))]
        if not len(v) and not len(rtu):
            continue
        rows.append(dict(sub_code=sc, bays=list(v["bay"]), rtu=bool(len(rtu)),
                         oldest=min(list(v["install_date"]) + list(rtu["install_date"]))))
    age = np.array([(W.ASOF - r["oldest"]).days / 365.25 for r in rows])
    order = np.argsort(-(age + rng.normal(0, 4, len(rows))), kind="stable")
    months = S.months_between(dt.date(2027, 1, 1), dt.date(2036, 1, 1))
    mw = np.array([S.SEASON[m.month - 1] for m in months])
    n_in = len(rows) - max(1, len(rows) // 7)
    picks = np.sort(rng.choice(len(months), n_in, replace=True, p=mw / mw.sum()))
    for rank, i in enumerate(order):
        m = months[picks[rank]] if rank < n_in else None
        rows[i]["plan_2026"] = m
        if m is not None and rng.random() < 0.2:
            m = max(dt.date(2027, 1, 1), S.month_add(m, -int(rng.integers(6, 19))))
        rows[i]["plan_2025"] = m
    return rows


def _schemes(rng, prog, extra):
    """stable scheme numbers: given when a scheme first entered the programme, so they do not follow the dates"""
    keys = sorted(list(prog["sub_code"][prog["retired_on"].isna()]) + [r["sub_code"] for r in extra])
    perm = rng.permutation(len(keys))
    return {k: "PRP-%03d" % (11 + 2 * int(p) + int(rng.integers(0, 2))) for k, p in zip(keys, perm)}


def programme_xlsx(path, rng, prog, subs, reg):
    wb = openpyxl.Workbook()
    names = dict(zip(subs["sub_code"], subs["sub_name"]))
    region = dict(zip(subs["sub_code"], subs["region"]))
    extra = other_schemes(rng, prog, reg)
    ids = _schemes(rng, prog, extra)
    pick = {k: int(rng.integers(0, 1000)) for k in ids}

    def rows_for(version):
        out = []
        for _, r in prog[prog["retired_on"].isna()].iterrows():
            sc, k = r.sub_code, pick[r.sub_code]
            whole = "%s-%s" % (r.first_bay, r.last_bay)
            when = r["plan_2026"] if version == 2026 else r["plan_2025"]
            if version == 2026 and isinstance(r["stage2_from"], str):
                cut = int(r["stage2_from"][1:])
                b1, b2 = "%s-F%02d" % (r.first_bay, cut - 1), "%s-%s" % (r["stage2_from"], r.last_bay)
                out.append((when, ids[sc], 1, sc, b1, SCOPE_S1[k % len(SCOPE_S1)] % b1))
                out.append((r["plan_2026_s2"], ids[sc], 2, sc, b2, SCOPE_S2[k % len(SCOPE_S2)] % b2))
            else:
                out.append((when, ids[sc] if when is not None else "", "", sc, whole, SCOPE_ONE[k % len(SCOPE_ONE)] % whole))
        for r in extra:
            sc, k = r["sub_code"], pick[r["sub_code"]]
            when = r["plan_2026"] if version == 2026 else r["plan_2025"]
            if r["bays"]:
                bays = _bay_list(r["bays"])
                scope = SCOPE_VPR220[k % len(SCOPE_VPR220)] % (bays + (" and RTU" if r["rtu"] else ""))
                if r["rtu"]:
                    bays += ", RTU"
            else:
                bays, scope = "RTU", "RTU replacement and SCADA cutover"
            out.append((when, ids[sc] if when is not None else "", "", sc, bays, scope))
        out.sort(key=lambda x: (x[0] or dt.date(2099, 1, 1), x[1] or "PRP-999", str(x[2]), x[3]))
        return out

    def sheet(ws, version, title):
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append([])
        hdr = ["Scheme", "Stage", "Substation code", "Substation", "Region", "Bays", "Scope", "Commissioning month"]
        ws.append(hdr)
        for j in range(1, len(hdr) + 1):
            ws.cell(row=3, column=j).font = Font(bold=True)
            ws.cell(row=3, column=j).fill = PatternFill("solid", fgColor="E3E9F2")
        for when, sid, stage, sc, bays, scope in rows_for(version):
            ws.append([sid or None, stage or None, sc, names[sc], region[sc], bays, scope,
                       dt.datetime(when.year, when.month, 1) if when is not None else "Not in programme period"])
            c = ws.cell(row=ws.max_row, column=8)
            if when is not None:
                c.number_format = "mmm yyyy"
            c.alignment = Alignment(horizontal="left")
        for j, w in enumerate([10, 7, 15, 20, 10, 16, 58, 22], 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = "A3:H%d" % ws.max_row

    ws = wb.active
    ws.title = "Approved 2026"
    sheet(ws, 2026, "Protection Renewal Programme 2026-2035, approved by the Board 19 March 2026 (resolution 2026-07)")
    ws2 = wb.create_sheet("2025 plan (superseded)")
    sheet(ws2, 2025, "Protection Renewal Programme, 2025 planning version (superseded by the approved 2026 programme)")
    ws3 = wb.create_sheet("Notes")
    notes = [
        "Commissioning month is the month the new protection for the listed bays goes into service. The relays in those "
        "bays come out of service in the same month.",
        "Large stations are built in two outage stages. Each stage lists the bays it replaces.",
        "Stations rebuilt under the 2022-2025 station renewals are complete and are not listed.",
        "Stations marked 'Not in programme period' keep their existing protection beyond December 2035.",
        "Scheme numbers are assigned when a scheme first enters the programme and do not change between versions.",
    ]
    for i, t in enumerate(notes, 1):
        ws3.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True)
    ws3.column_dimensions["A"].width = 110
    wb.save(path)


def _bay_list(bays):
    """F01, F03-F05 style list of feeder bays"""
    nums = sorted(int(b[1:]) for b in bays)
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        out.append("F%02d" % nums[i] if i == j else "F%02d-F%02d" % (nums[i], nums[j]))
        i = j + 1
    return ", ".join(out)


def dictionary_txt(path):
    txt = """TARROCK HYDRO DISTRIBUTION
EAM reporting service - monthly P&C asset and work extract (job PC_ASSET_WO_MONTHLY)
Run 2026-10-01 06:10, data to close of business 2026-09-30

General
  The EAM went live on 1 January 2021, replacing the legacy maintenance system. Assets in service at go-live,
  and the rotating items fitted to them, were loaded from the legacy lists. Work orders raised before go-live
  were archived with the legacy system and are not in EAM. Assets retired before go-live are not in EAM.
  Operating regions were reorganised on 1 April 2024 (four regions to three); crews were renamed at the same time.
  Dates are local (Eastern). Blank means no value.

protection_asset_register_2026-09-30.csv  (one row per asset)
  asset_id          EAM asset number (shared numbering with all asset classes)
  model, manufacturer, function
  sub_code          substation code, see substations.csv
  bay               bay or position: Fnn feeder, Tn transformer, Bn bus zone, RTU
  serial_no         manufacturer serial number
  install_date      date placed in service (legacy list date for assets loaded at go-live)
  status            IN SERVICE or RETIRED
  retired_date      date taken out of service
  region            region on the asset record when it was created (not updated by the 2024 reorganisation)

psu_rotating_item_history_2021-01_2026-09.csv  (item VPR340-PSU-02, one row per module per installation)
  serial_no         module serial number; a module repaired by the manufacturer keeps its serial
  asset_id          relay the module is or was fitted to
  fitted_date       date fitted to that relay (legacy fitted date for modules loaded at go-live)
  fitting_wo        work order the module was fitted under (blank for go-live loads)
  removed_date      date removed (blank while fitted)
  removal_reason    FAILED, or ASSET RETIRED when the relay was removed by a station rebuild
  removal_wo        work order the module was removed under
  record_source     GO-LIVE LOAD or EAM

eam_work_orders_protection_2021-01_2026-09.csv  (one row per work order)
  wo_number         work-order number (shared numbering with lines and metering work)
  reported_date     date raised
  completed_date    date completed (blank while open or cancelled)
  work_type         PM periodic test, CM corrective, EM emergency corrective, INSP inspection,
                    MOD modification (settings, firmware, device replacement or station rebuild cutover),
                    LOG logistics (shipping and returns)
  status            INPRG in progress, COMP completed, CLOSE closed after supervisor review, CANC cancelled
  asset_id          blank for station-level inspections
  sub_code, region  region at the date raised
  failed_component  recorded by the crew on corrective and logistics work: PSU power supply, CPU processor,
                    IO input/output card, DISP display, COMMS communications card (blank when no fault was found)
  description       crew text
  crew              crew lead's crew
  labour_hours      hours booked, including travel

substations.csv
  sub_code, sub_name, region (from 1 April 2024), legacy_region (before 1 April 2024),
  station_type (TS transformer station, DS distribution station)

spares_stock_2026-09-30.csv  (stores count at close of business 2026-09-30)
  stock_state       SERVICEABLE on the shelf; QUARANTINE awaiting bench test (returned from site or recovered
                    from a station rebuild); AT REPAIR at the manufacturer under an open RMA
  unit_cost_cad     moving average cost
"""
    open(path, "w").write(txt)
