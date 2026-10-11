"""Assertions on the shipped pack: input gates, the answer, the ladder, bin margins, rounding edges."""
import datetime as dt
import json
import os
import re
import statistics
import sys
import zipfile

import docx
import openpyxl
import pandas as pd
from pypdf import PdfReader

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import estimators as E
import world as W

TASK = os.path.dirname(HERE)
T = os.path.join(TASK, "target")
n_ok = 0
LOTS = 13


def ok(cond, msg):
    global n_ok
    if not cond:
        raise AssertionError(msg)
    n_ok += 1


def edge(v):
    return abs((v * 10) % 1 - 0.5)


def main():
    files = sorted(os.listdir(T))
    ok(len(files) >= 10, "10 or more input files")
    ok(len({f.rsplit(".", 1)[1] for f in files}) >= 3, "3 or more formats")
    wo = pd.read_csv(os.path.join(T, E.F["wo"]), dtype=str)
    ok(len(wo) >= 25000, "work-order spine has 25,000 or more rows (%d)" % len(wo))
    ok(wo["reported_date"].min() >= "2021-01-01", "no work order before the EAM go-live")
    meta = json.load(open(os.path.join(TASK, "metadata.json")))
    ok(len(meta["distractors"]) >= 2 and all(d in files for d in meta["distractors"]), "two or more distractors on disk")
    for f in files:
        ok(not any(w in f.lower() for w in ("answer", "golden", "trap", "clean", "_v2", "test_")), "file name %s carries no author noun" % f)

    s = E.solve(T)
    ok(s["lots"] == LOTS, "answer is %d lots (%d)" % (LOTS, s["lots"]))
    ok(s["cover"] >= W.COVER > s["settings"][s["lots"] - 1]["cover"], "the answer clears 95 per cent and one lot fewer does not")
    ok(min(W.SETTINGS) <= s["lots"] - 3 and max(W.SETTINGS) >= s["lots"] + 5, "the asked settings run well either side of the answer")
    lo, hi = s["lam"], s["lam"]
    while E.lots_needed(lo, s["avail"]) == s["lots"]:
        lo -= 0.05
    while E.lots_needed(hi, s["avail"]) == s["lots"]:
        hi += 0.05
    ok((s["lam"] - lo) / s["lam"] >= 0.025 and (hi - s["lam"]) / s["lam"] >= 0.025,
       "expected failures sit at least 2.5 per cent inside the answer's lot bin (%.1f < %.1f < %.1f)" % (lo, s["lam"], hi))
    vals = [s["lam"]] + list(s["by_year"].values()) + list(s["by_region"].values()) + [
        100 * v["cover"] for v in s["settings"].values() if 0.05 < 100 * v["cover"] < 99.95]
    for v in vals:
        ok(edge(v) >= 0.04, "figure %.4f clear of its one-decimal rounding edge" % v)
    ok(abs(sum(s["by_year"].values()) - s["lam"]) < 1e-6, "years add to the total")
    ok(abs(sum(s["by_region"].values()) - s["lam"]) < 1e-6, "regions add to the total")
    ok(list(s["in_service"].values()) == sorted(s["in_service"].values(), reverse=True), "supported relays fall year on year")

    r, avail = E.rungs(T)
    ans = r["ANSWER"][0]
    for k, (v, l) in r.items():
        if k == "ANSWER":
            continue
        ok(l != s["lots"], "rung %s lands on %d lots, not the answer" % (k, l))
    ok(r["window_no_renewal"][1] >= s["lots"] + 3, "keeping today's module in every month is at least three lots high")
    ok(r["window_no_renewal"][0] >= 1.25 * ans, "and at least 25 per cent high")
    ok(r["install_renewal"][0] <= 0.88 * ans, "exposure from fitting rather than the go-live is at least 12 per cent low")
    ok(r["relay_age_window_no_renewal"][0] <= 0.88 * ans, "relay age in place of module age is at least 12 per cent low")
    ok(r["trailing_rate"][0] <= 0.80 * ans, "the trailing twelve-month rate is at least 20 per cent low")
    # staged schemes and the superseded plan show in the supported-relay column
    reg0, wo0, hist0, subs0, stock0, plan0 = E.load(T)
    fl0 = E.fleet(reg0)
    in25 = E.in_service_jan1(fl0, plan0["2025 plan (superseded)"])
    ok(sum(in25[y] != s["in_service"][y] for y in in25) >= 6, "the 2025 plan misstates the supported relays in most years")
    p26 = plan0["Approved 2026"]
    first = {sub: min((w or E.month_add(W.PROGRAMME_END, 1)) for _, w in rows) for sub, rows in p26.items()}
    live = fl0[fl0["status"] == "IN SERVICE"]
    by_station = {y: int(sum(1 for sub in live["sub_code"] if sub in first and first[sub] > dt.date(y, 1, 1))) for y in range(2027, 2036)}
    ok(sum(by_station[y] != s["in_service"][y] for y in by_station) >= 5, "reading a staged scheme by station misstates the supported relays")
    # every module counted in the rates is in the history the procedure names, and the failures tie to it
    f = E.failures(wo0, fl0, hist0)
    ok(len(f) == s["n_fail"] and f["fitted"].notna().all(), "every failure ties to the module it removed")
    ok((hist0["removal_reason"] == "FAILED").sum() >= len(f), "the history records a removal for every failure")
    realism(reg0, wo0, hist0, subs0, plan0, f)
    print("CHECKS PASS", n_ok)



def _text(path):
    """the words a reader sees in a shipped file"""
    if path.endswith(".docx"):
        d = docx.Document(path)
        out = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                out += [c.text for c in row.cells]
        for sec in d.sections:
            out += [p.text for p in sec.footer.paragraphs]
        return "\n".join(out)
    if path.endswith(".pdf"):
        return "\n".join(pg.extract_text() for pg in PdfReader(path).pages)
    if path.endswith(".xlsx"):
        wb = openpyxl.load_workbook(path, data_only=True)
        return "\n".join(" ".join(str(v) for v in row if v is not None) for ws in wb for row in ws.iter_rows(values_only=True)) + \
            "\n" + " ".join(wb.sheetnames)
    return open(path, encoding="utf-8").read()


def realism(reg, wo, hist, subs, plan, f):
    """the operating texture the pack has to carry, each piece asserted so a later retune cannot strip it out"""
    d = lambda x: dt.date.fromisoformat(x) if x else None
    wo_by = wo.set_index("wo_number")
    # failure path: some spares fitted on a follow-up work order raised the same or the next working day, numbered
    # after the failure, carrying no failed component, and fitted no later than the first day of the next month
    eam = hist[hist["record_source"] == "EAM"]
    sep = 0
    for a, fit, fwo in zip(eam["asset_id"], eam["fitted"], eam["fitting_wo"]):
        prev = hist[(hist["asset_id"] == a) & hist["removed"].map(lambda x: x is not None) & (hist["removal_reason"] == "FAILED")]
        prev = prev[prev["removed"] <= fit].sort_values("removed").iloc[-1]
        nxt = E.month_add(dt.date(prev["removed"].year, prev["removed"].month, 1), 1)
        ok(prev["removed"] <= fit <= nxt, "replacement on %s fitted between the removal and the next first of the month" % a)
        if fwo != prev["removal_wo"]:
            sep += 1
            w = wo_by.loc[fwo]
            ok(fwo > prev["removal_wo"] and w["work_type"] == "CM" and w["failed_component"] == "" and
               (d(w["reported_date"]) - prev["removed"]).days <= 4, "follow-up fitting %s raised after its failure, no component" % fwo)
    ok(sep >= 15, "%d spares fitted on a separate follow-up work order" % sep)
    ok(len(f) == (hist["removal_reason"] == "FAILED").sum(), "every 5.1 failure is one FAILED removal")
    # RMA register: from the 2015 service bulletin, boxes of two, references blank or mistyped, skewed receipt lags,
    # a module removed before go-live received in January 2021
    wb = openpyxl.load_workbook(os.path.join(T, "veldtek_rma_log_vpr340_psu.xlsx"))
    rows = list(wb.active.iter_rows(min_row=2, values_only=True))
    rec = [r[5].date() for r in rows]
    ok(min(rec) < dt.date(2015, 12, 31) and sum(x < E.month_add(dt.date(2021, 1, 1), 0) for x in rec) >= 50,
       "the RMA register runs from 2015 (%d receipts before go-live)" % sum(x < dt.date(2021, 1, 1) for x in rec))
    ok(sum("SB-15-02" in str(r[6]) for r in rows) >= 30, "the service-bulletin returns are in the register")
    created = re.search(r"<dcterms:created[^>]*>(\d{4}-\d{2}-\d{2})", zipfile.ZipFile(
        os.path.join(T, "veldtek_rma_log_vpr340_psu.xlsx")).read("docProps/core.xml").decode()).group(1)
    ok(created < min(rec).isoformat(), "the register workbook was created before its first receipt")
    log = wo[(wo["work_type"] == "LOG") & (wo["failed_component"] == "PSU")]
    refs = [str(r[4]) for r in rows if r[4]]
    pairs = sum(1 for w in set(log["wo_number"]) if refs.count(w) == 2)
    ok(pairs >= 5, "%d logistics work orders ship two modules in one box" % pairs)
    blank = sum(1 for r in rows if not r[4] and r[5].year >= 2021)
    odd = sum(1 for r in rows if r[4] and r[5].year >= 2021 and str(r[4]) not in set(log["wo_number"]))
    ok(blank >= 3 and odd >= 3, "references blank (%d) or not a logistics work order (%d) after go-live" % (blank, odd))
    serials = set(hist["serial_no"])
    jan = [r for r in rows if r[5].year == 2021 and r[5].month == 1 and r[2] not in serials]
    ok(len(jan) >= 1, "a module removed before go-live received in January 2021")
    removed = {}
    for s_, rm, why in zip(hist["serial_no"], hist["removed"], hist["removal_reason"]):
        if why == "FAILED":
            removed.setdefault(s_, []).append(rm)
    lag = []
    for r in rows:
        prior = [x for x in removed.get(r[2], []) if x <= r[5].date()]
        if prior:
            lag.append((r[5].date() - max(prior)).days)
    ok(statistics.median(lag) <= 25 and max(lag) >= 90 and statistics.mean(lag) > statistics.median(lag),
       "removal-to-receipt days skewed (median %s, mean %.0f, max %s)" % (statistics.median(lag), statistics.mean(lag), max(lag)))
    turn = [(r[7] - r[5]).days for r in rows if r[7] is not None]
    ok(min(turn) <= 15 and max(turn) >= 80, "repair turnaround from %d to %d days" % (min(turn), max(turn)))
    # master data: year-only legacy dates as 1 July, units held in stores before installation, the feeder identity
    ok(sum(str(x).endswith("-07-01") for x in reg["install_date"]) >= 30, "legacy install dates held as 1 July")
    ok(sum(1 for x, s_ in zip(hist["fitted_date"], hist["record_source"]) if s_ == "GO-LIVE LOAD" and x.endswith("-07-01")) >= 5,
       "legacy fitted dates held as 1 July")
    pat = {"VPR-340": r"V340-(\d\d)-", "VPR-220": r"V220-(\d\d)-", "NX-9": r"N9-(\d\d)-", "NX-5": r"N5-(\d\d)-",
           "AT-7": r"AW7(\d\d)", "KDR-30": r"K30(\d\d)", "RT-400": r"AW4(\d\d)", "RT-600": r"AW6(\d\d)"}
    gap = []
    for mdl, sn, inst in zip(reg["model"], reg["serial_no"], reg["install_date"]):
        yy = int(re.match(pat[mdl], sn).group(1))
        gap.append(int(str(inst)[:4]) - (1900 + yy if yy > 50 else 2000 + yy))
    ok(min(gap) >= 0 and max(gap) <= 3 and sum(g >= 2 for g in gap) >= 20,
       "serial year 0 to 3 years before installation (%d at 2 or 3)" % sum(g >= 2 for g in gap))
    trips = pd.read_csv(os.path.join(T, "feeder_relay_operations_2025.csv"))
    brk = int((trips["trip_operations"] != trips["reclose_attempts"] + trips["lockouts"]).sum())
    ok(50 <= brk <= 0.25 * len(trips), "%d feeders log trips outside the reclosing identity" % brk)
    # program: spare bays and VPR-220 feeders inside scheme bay lists, a 2025 plan that differs, every live VPR-340 bay
    # in exactly one approved row
    p26, p25 = plan["Approved 2026"], plan["2025 plan (superseded)"]
    live = reg[reg["status"] == "IN SERVICE"]
    occupied = set(zip(reg["sub_code"], reg["bay"]))
    spare = mixed = 0
    for sub, rws in p26.items():
        listed = {b for rngs, _ in rws for lo, hi in rngs for b in range(lo, hi + 1)}
        spare += sum((sub, "F%02d" % b) not in occupied for b in listed)
        models = set(live[(live["sub_code"] == sub) & live["bay"].isin({"F%02d" % b for b in listed})]["model"])
        mixed += {"VPR-340", "VPR-220"} <= models
    ok(spare >= 3 and mixed >= 3, "approved schemes take in %d spare bays and %d stations mixing VPR-220 and VPR-340" % (spare, mixed))
    v340 = live[live["model"] == "VPR-340"]
    for sub, bay in zip(v340["sub_code"], v340["bay"]):
        n = sum(any(lo <= int(bay[1:]) <= hi for lo, hi in rngs) for rngs, _ in p26.get(sub, []))
        ok(n == 1, "VPR-340 %s %s sits in one approved row" % (sub, bay))
    ok(len(set(p26) ^ set(p25)) >= 10, "the 2025 plan lists a different station set")
    wbp = openpyxl.load_workbook(os.path.join(T, E.F["prog"]))
    sc26 = {r[2]: r[6] for r in wbp["Approved 2026"].iter_rows(min_row=4, values_only=True) if r[2] and r[1] is None}
    sc25 = {r[2]: r[6] for r in wbp["2025 plan (superseded)"].iter_rows(min_row=4, values_only=True) if r[2]}
    both = set(sc26) & set(sc25)
    ok(0.2 <= sum(sc26[k] != sc25[k] for k in both) / len(both) <= 0.8, "scope text differs between versions in places")
    notes = " ".join(str(r[0]) for r in wbp["Notes"].iter_rows(values_only=True))
    ok("Each stage replaces the bays it lists" in notes and
       "out of service from the first day of the commissioning month" in notes, "the program notes pin staging and the outage day")
    ok("keep their existing protection" not in notes, "no note written to an analyst")
    # maintenance distractor derived from the PM work orders, with formulas
    import distract as Q
    wo_d = wo.copy()
    wo_d["reported_date"] = wo_d["reported_date"].map(d)
    wo_d["completed_date"] = wo_d["completed_date"].map(d)
    cnt = Q.compliance_counts(reg, subs, wo_d)
    wc = openpyxl.load_workbook(os.path.join(T, "pc_relay_maintenance_compliance_2025.xlsx"))
    ws = wc.active
    body = [r for r in ws.iter_rows(min_row=4, values_only=True) if r[0] and r[0] != "All models"]
    ok(len(body) == len(cnt) and all(list(cnt[(r[0], r[1])]) == list(r[2:6]) for r in body),
       "compliance counts recompute from the 2025 PM work orders")
    ok(all(str(c.value).startswith("=") for c in ws["G"][3:]), "on-time per cent is a formula")
    pm = wo[(wo["work_type"] == "PM") & (wo["status"] != "CANC") & wo["description"].str.contains("relay test|RELAY MAINT", case=False)]
    kind = dict(zip(subs["sub_code"], subs["station_type"]))
    dev = []
    for (a, sub), g in pm.groupby(["asset_id", "sub_code"]):
        ds = sorted(d(x) for x in g["reported_date"])
        dev += [(y - x).days - (730 if kind[sub] == "TS" else 1095) for x, y in zip(ds, ds[1:])]
    ok(statistics.pstdev(dev) > 15 and sum(x > 60 for x in dev) >= 20,
       "relay test intervals slip past the nominal interval (sd %.0f days, %d over 60)" % (statistics.pstdev(dev), sum(x > 60 for x in dev)))
    # crew text and the extract's make-up
    fdesc = wo[wo["wo_number"].isin(set(hist["removal_wo"][hist["removal_reason"] == "FAILED"]))]["description"]
    ok(fdesc.nunique() >= 0.95 * len(fdesc), "failure write-ups are the crews' own (%d distinct of %d)" % (fdesc.nunique(), len(fdesc)))
    ok(wo["description"].nunique() >= 5000, "%d distinct descriptions" % wo["description"].nunique())
    ok(not wo["description"].str.contains(r"\bu/s\b|healthy", case=False).any(), "no UK crew terms")
    back = {}
    for r in rows:
        if r[7] is not None:
            back.setdefault(r[2], []).append(r[7].date())
    rep = 0
    for s_, fit, fw in zip(hist["serial_no"], hist["fitted"], hist["fitting_wo"]):
        if fw and any(x < fit for x in back.get(s_, [])):
            rep += 1
            t = wo_by.loc[fw, "description"].lower()
            ok(not re.search(r"\bnew (module|psu|unit|power supply|p/s)", t), "repaired %s not written up as new" % s_)
    ok(rep >= 10, "%d repaired modules fitted again" % rep)
    ok(len(wo) >= 25000 and (wo["work_type"] == "INSP").mean() <= 0.25 and (wo["asset_id"] == "").mean() <= 0.2,
       "extract over 25,000 rows, inspections %.0f%%, blank asset %.0f%%" % (100 * (wo["work_type"] == "INSP").mean(),
                                                                            100 * (wo["asset_id"] == "").mean()))
    other = set(wo["asset_id"]) - set(reg["asset_id"]) - {""}
    ok(len(other) >= 1000 and not other & set(reg["asset_id"]), "%d station DC, breaker and VT assets carry their own numbers" % len(other))
    # documents: Canadian spelling, the letter date, repair terms, the procedure's filing rule and definitions
    for fn in sorted(os.listdir(T)):
        txt = _text(os.path.join(T, fn))
        ok(not re.search(r"programme|organis|reorganis|minimis|prioritis|analys(e|ing)\b", txt, re.I), "Canadian spelling in %s" % fn)
    eol = _text(os.path.join(T, "veldtek_vpr340_eol_notice.pdf"))
    ok("August 20, 2026" in eol and "electrolytic capacitors and fan" in eol.replace("\n", " "), "the notice's date and repair terms")
    proc = _text(os.path.join(T, "AM-PC-014_obsolete_spares_final_buy.docx"))
    ok("by year" not in proc and "Monthly failure rate" in proc and "Module age" in proc, "8.3 files, section 3 defines")
    ok(not re.search(r"ageing|aging|renewed|starts again|replacement module", proc, re.I), "no sentence on ageing or renewal")
    dic = _text(os.path.join(T, "eam_extract_dictionary.txt"))
    for col in ["item_code", "item_no", "item_description", "location", "bin", "stock_state", "uom", "quantity", "unit_cost_cad",
                "1 July", "KALLISON"]:
        ok(col in dic, "dictionary documents %s" % col)
    # containers: each workbook saved by its own Excel, Word's ids irregular, the PDF stamped to the second
    geo = set()
    for fn in ("protection_renewal_program_2026.xlsx", "veldtek_rma_log_vpr340_psu.xlsx", "pc_relay_maintenance_compliance_2025.xlsx"):
        x = zipfile.ZipFile(os.path.join(T, fn)).read("xl/workbook.xml").decode()
        geo.add((re.search(r'rupBuild="(\d+)"', x).group(1), re.search(r"<workbookView [^>]*", x).group(0)[:80]))
    ok(len(geo) == 3, "three workbooks, three Excel builds and windows")
    x = zipfile.ZipFile(os.path.join(T, "AM-PC-014_obsolete_spares_final_buy.docx")).read("word/document.xml").decode()
    tids = re.findall(r'w14:textId="([0-9A-F]{8})"', x)
    ok("77777777" not in tids and len(set(tids)) == len(tids), "text ids are Word's own")
    runs = [len(list(g)) for _, g in __import__("itertools").groupby(re.findall(r'<w:p [^>]*w:rsidR="([0-9A-F]+)"', x))]
    ok(len(set(runs)) >= 5, "revision sessions run in blocks of %d different lengths" % len(set(runs)))
    info = PdfReader(os.path.join(T, "veldtek_vpr340_eol_notice.pdf")).metadata
    ok(not info["/CreationDate"][14:16] == "00", "PDF creation stamped to the second")


if __name__ == "__main__":
    main()
