"""The golden deliverables: the final-buy workbook and the memo to the Director, from the shipped pack."""
import datetime as dt
import io
import os
import sys
from zoneinfo import ZoneInfo

import docx
import matplotlib
import openpyxl
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

matplotlib.use("Agg")
import matplotlib.pyplot as plt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import estimators as E
import postprocess as P
import world as W

TASK = os.path.dirname(HERE)
PREPARED = dt.datetime(2026, 10, 7, 15, 40)
TZ = ZoneInfo("America/Toronto")
INK, ACCENT, MUTED, BAR = "#1F2A33", "#2F6B55", "#8A979F", "#9DBFB0"


def _hdr(ws, row, vals, widths):
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=j, value=v)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F6B55")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    ws.row_dimensions[row].height = 30


def workbook(s, path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Buy"
    ws["A1"] = "VPR-340 power supply module (VPR340-PSU-02): last-time buy"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "P&C Asset Planning, prepared 7 October 2026 for the Director, Asset Management. Sized under AM-PC-014 rev 4."
    ws["A2"].font = Font(italic=True, color="5A6A72")
    rows = [
        ("Lots to order (25 modules per lot)", s["lots"], "0"),
        ("Modules to order", s["modules"], "#,##0"),
        ("Cost at the last-time-buy price (CAD)", s["cost"], '"$"#,##0'),
        ("Expected failures, 1 Oct 2026 to end of support", round(s["lam"], 1), "#,##0.0"),
        ("Available modules at the 30 Sep 2026 count (serviceable plus at Veldtek repair)", s["avail"], "0"),
        ("Probability the stock covers every failure to end of support", round(100 * s["cover"], 1), '0.0"%"'),
    ]
    for i, (k, v, fmt) in enumerate(rows, 4):
        ws.cell(row=i, column=1, value=k)
        c = ws.cell(row=i, column=2, value=v)
        c.number_format = fmt
        c.alignment = Alignment(horizontal="right")
    ws.cell(row=4, column=1).font = Font(bold=True)
    ws.cell(row=4, column=2).font = Font(bold=True)
    ws.column_dimensions["A"].width = 74
    ws.column_dimensions["B"].width = 16

    ws = wb.create_sheet("Lot settings")
    _hdr(ws, 1, ["Lots ordered", "Modules ordered", "Modules available with the buy", "Cover probability (%)",
                 "Cost (CAD)", "Meets 95%"], [12, 15, 18, 16, 15, 11])
    for i, n in enumerate(W.SETTINGS, 2):
        st = s["settings"][n]
        vals = [n, n * W.LOT, s["avail"] + n * W.LOT, round(100 * st["cover"], 1), round(st["cost"]), "yes" if st["cover"] >= W.COVER else "no"]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.number_format = ["0", "#,##0", "#,##0", "0.0", '"$"#,##0', "@"][j - 1]
        if n == s["lots"]:
            for j in range(1, 7):
                ws.cell(row=i, column=j).font = Font(bold=True)
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Outlook by year")
    _hdr(ws, 1, ["Year", "Expected PSU failures", "Relays supported on 1 January"], [14, 20, 22])
    for i, (y, v) in enumerate(s["by_year"].items(), 2):
        ws.cell(row=i, column=1, value="2026 (Oct-Dec)" if y == 2026 else y)
        ws.cell(row=i, column=2, value=round(v, 1)).number_format = "0.0"
        if y in s["in_service"]:
            ws.cell(row=i, column=3, value=s["in_service"][y]).number_format = "#,##0"
    n = len(s["by_year"]) + 2
    ws.cell(row=n, column=1, value="Total").font = Font(bold=True)
    c = ws.cell(row=n, column=2, value=round(s["lam"], 1))
    c.number_format = "0.0"
    c.font = Font(bold=True)
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("By region")
    _hdr(ws, 1, ["Operating region (from April 2024)", "Expected PSU failures to end of support"], [30, 22])
    for i, r in enumerate(["North", "Central", "South"], 2):
        ws.cell(row=i, column=1, value=r)
        ws.cell(row=i, column=2, value=round(s["by_region"][r], 1)).number_format = "0.0"

    ws = wb.create_sheet("Failure rates")
    _hdr(ws, 1, ["Module age band (years)", "PSU failures Jan 2021 to Sep 2026", "Module-years in service Jan 2021 to Sep 2026",
                 "Failures per module-year", "Expected module-years to end of support", "Expected failures"], [16, 18, 20, 16, 20, 14])
    for i, b in enumerate(W.BAND_LABELS, 2):
        k = i - 2
        ws.cell(row=i, column=1, value=b)
        ws.cell(row=i, column=2, value=int(s["counts"][k]))
        ws.cell(row=i, column=3, value=round(s["exposure"][k] / 12.0, 1)).number_format = "#,##0.0"
        ws.cell(row=i, column=4, value=round(float(s["rates"][k]), 4)).number_format = "0.0000"
        ws.cell(row=i, column=5, value=round(s["occupancy"][k] / 12.0, 1)).number_format = "#,##0.0"
        ws.cell(row=i, column=6, value=round(float(s["rates"][k] * s["occupancy"][k] / 12.0), 1)).number_format = "0.0"
    n = len(W.BAND_LABELS) + 2
    ws.cell(row=n, column=1, value="Total").font = Font(bold=True)
    ws.cell(row=n, column=2, value=int(s["counts"].sum())).font = Font(bold=True)
    c = ws.cell(row=n, column=6, value=round(s["lam"], 1))
    c.number_format = "0.0"
    c.font = Font(bold=True)
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Notes")
    notes = [
        "Failures: corrective (CM, EM) work orders on VPR-340 relays with failed component PSU, status COMP or CLOSE (AM-PC-014 5.1). "
        "Each failure takes the age of the module it removed, from that module's fitted date in the rotating-item history (5.2).",
        "Rates: module-years are counted from the January 2021 go-live, when the work-order history begins; a module fitted earlier "
        "enters the count at go-live at the age it had then (5.3).",
        "Expected module-years: each supported relay's present module is aged month by month. A module that fails is replaced from "
        "stock by a new module, which counts in the 0-9 band from the following month, so each month's module-years are split between "
        "the present module (by the chance it has not yet failed) and its replacement (6.1, 6.2).",
        "Support ends in the commissioning month of the scheme, or the stage of a staged scheme, that replaces the relay's bay in the "
        "approved 2026 programme; relays at stations with no scheme in the period are supported to December 2035 (4.1, 4.2).",
        "Available modules: %d serviceable at Central Stores, North and South depots plus %d at Veldtek under open RMA; %d quarantined "
        "modules excluded (7)." % (s["serviceable"], s["at_repair"], s["quarantine"]),
        "Cover is the Poisson probability that available modules plus the buy meet all failures to end of support (8.1, 8.2).",
    ]
    for i, t in enumerate(notes, 1):
        ws.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 120
    wb.save(path)


def chart_png(s):
    years = list(s["by_year"].keys())
    vals = [s["by_year"][y] for y in years]
    labels = ["Q4 2026" if y == 2026 else str(y) for y in years]
    fig, ax = plt.subplots(figsize=(7.2, 2.7), dpi=200)
    ax.bar(labels, vals, color=[BAR] * len(vals), width=0.62, edgecolor="none")
    for x, v in zip(labels, vals):
        ax.text(x, v + 1.2, "%.1f" % v, ha="center", va="bottom", fontsize=7.5, color=INK)
    ax.set_ylabel("Expected PSU failures", fontsize=8.5, color=INK)
    ax.set_ylim(0, max(vals) * 1.18)
    ax.spines[["top", "right"]].set_visible(False)
    ax.tick_params(labelsize=8, colors=INK)
    ax.grid(axis="y", color="#E3E8EA", linewidth=0.6)
    ax.set_axisbelow(True)
    ax.set_title("Expected VPR-340 PSU failures by year until the renewal programme replaces each relay",
                 fontsize=9, color=INK, loc="left")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    buf.seek(0)
    return buf


def memo(s, path):
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for sec in d.sections:
        sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
        sec.left_margin = sec.right_margin = Cm(2.3)
        sec.top_margin = Cm(1.4)
        sec.bottom_margin = Cm(1.4)
    cp = d.core_properties
    cp.author = "Rachael Villanueva"
    cp.last_modified_by = "Rachael Villanueva"
    cp.title = "VPR-340 PSU last-time buy"
    cp.created = cp.modified = PREPARED
    cp.revision = 2
    cp.comments = cp.keywords = cp.subject = cp.category = ""

    def para(t, bold=False, size=None, after=5, italic=False):
        p = d.add_paragraph()
        r = p.add_run(t)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(after)
        return p

    para("Tarrock Hydro Distribution  |  P&C Asset Planning", size=8.5, italic=True, after=2)
    para("MEMO", bold=True, size=12, after=2)
    for k, v in [("To", "Angel Baird, Director, Asset Management"), ("From", "Rachael Villanueva, Manager, P&C Asset Planning"),
                 ("Date", "7 October 2026"), ("Re", "Last-time buy of VPR-340 power supply modules (Veldtek EOL-2026-07)")]:
        p = d.add_paragraph()
        r1 = p.add_run(k + ":\t")
        r1.bold = True
        p.add_run(v)
        p.paragraph_format.space_after = Pt(0)
    para("", after=4)
    para("Order %d lots (%d modules) of VPR340-PSU-02 for CAD %s before Veldtek's 15 December deadline."
         % (s["lots"], s["modules"], "{:,.0f}".format(s["cost"])), bold=True, size=11.5, after=6)
    para("With the %d modules we can use today, %d lots give a %.1f per cent chance of covering every power supply failure "
         "until the renewal programme reaches each relay, against the 95 per cent AM-PC-014 asks for; %d lots would give %.1f per "
         "cent. We expect %.1f failures between October 2026 and the end of support, most of them in 2027 to 2031 while the fleet is "
         "still large."
         % (s["avail"], s["lots"], 100 * s["cover"], s["lots"] - 1, 100 * s["settings"][s["lots"] - 1]["cover"], s["lam"]), after=6)
    para("How the figure is built. Rates are set for the AM-PC-014 module age bands from the PSU failures in the EAM history since "
         "January 2021, over the module-years each band has been in service since then; the 25-and-over band fails about a "
         "quarter of its modules a year. Each relay's present module is then aged month by month until the programme reaches "
         "its bay. A module that fails comes out and its replacement from this buy starts again in the 0-9 band, where failures "
         "are rare, so the old module's rate only applies for as long as it is likely to still be there.", after=6)
    pic = chart_png(s)
    d.add_picture(pic, width=Cm(15.0))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para("Expected failures by year (one decimal) from the workbook's Outlook sheet; Q4 2026 is October to December.", size=8.5, italic=True, after=8)
    t = d.add_table(rows=1, cols=4)
    t.style = "Table Grid"

    def cell(c, text, bold=False, right=False):
        c.text = ""
        r = c.paragraphs[0].add_run(text)
        if bold:
            r.bold = True
        r.font.size = Pt(9)
        c.paragraphs[0].paragraph_format.space_after = Pt(0)
        if right:
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    for j, h in enumerate(["Lots ordered", "Modules available with the buy", "Cover (%)", "Cost (CAD)"]):
        cell(t.rows[0].cells[j], h, bold=True, right=j > 0)
    for n in (s["lots"] - 1, s["lots"], s["lots"] + 1):
        stg = s["settings"][n]
        row = t.add_row().cells
        b = n == s["lots"]
        cell(row[0], str(n), bold=b)
        cell(row[1], "{:,}".format(s["avail"] + n * W.LOT), bold=b, right=True)
        cell(row[2], "%.1f" % (100 * stg["cover"]), bold=b, right=True)
        cell(row[3], "{:,.0f}".format(stg["cost"]), bold=b, right=True)
    para("", after=4)
    para("Available modules are the %d serviceable at Central Stores and the two depots plus the %d at Veldtek under open RMA, "
         "which the notice says will come back repaired; the %d in quarantine are left out until they pass bench test. The "
         "workbook vpr340_psu_final_buy.xlsx carries every order from %d to %d lots, the outlook by year and region, and the "
         "band rates." % (s["serviceable"], s["at_repair"], s["quarantine"], min(W.SETTINGS), max(W.SETTINGS)), size=9.5, after=6)
    para("Signed for order:  ______________________   Angel Baird, Director, Asset Management", after=0)
    d.save(path)


def build(target=None, out=None):
    target = target or os.path.join(TASK, "target")
    out = out or os.path.join(TASK, "golden")
    os.makedirs(out, exist_ok=True)
    s = E.solve(target)
    wbp = os.path.join(out, "vpr340_psu_final_buy.xlsx")
    workbook(s, wbp)
    utc = lambda t: dt.datetime(*t, tzinfo=TZ).astimezone(dt.timezone.utc).timetuple()[:6]
    P.normalize_xlsx(wbp, "Rachael Villanueva", "Rachael Villanueva", utc((2026, 10, 6, 13, 55, 18)),
                     utc((2026, 10, 7, 15, 38, 46)), seed="vpr340_psu_final_buy.xlsx")
    mp = os.path.join(out, "psu_last_time_buy_memo.docx")
    memo(s, mp)
    pages = P.normalize_docx(mp, "Rachael Villanueva", "Rachael Villanueva", utc((2026, 10, 6, 14, 2, 0)),
                             utc((2026, 10, 7, 15, 40, 0)), revision=6, total_time=74, title="VPR-340 PSU last-time buy",
                             seed="memo-psu")
    assert pages == 1, "the memo runs to %d pages" % pages
    for f in (wbp, mp):
        ts = PREPARED.replace(tzinfo=TZ).timestamp()
        os.utime(f, (ts, ts))
    return s


if __name__ == "__main__":
    s = build()
    print("lots", s["lots"], "lam", round(s["lam"], 1), "cover", round(s["cover"], 4))
