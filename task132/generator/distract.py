"""Files a protection team keeps that look relevant to the buy and do not size it."""
import datetime as dt

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill

Y0, Y1 = dt.date(2025, 1, 1), dt.date(2026, 1, 1)
SHOTS = {0: "OFF", 1: "1 SHOT", 2: "2 SHOT"}


def _in_2025(reg):
    ret = [r if isinstance(r, dt.date) else None for r in reg["retired_date"]]
    return reg[[i < Y1 and (r is None or r > Y0) for i, r in zip(reg["install_date"], ret)]]


def trip_operations_csv(path, rng, reg):
    """SCADA-logged protection operations per feeder relay or recloser control in 2025. Each feeder has its own
    reclosing: off (mostly cable feeders), one shot or two shots (rural overhead feeders and recloser controls).
    A fault trips the breaker; each reclose either holds or is followed by another trip, and a fault the last shot
    does not clear ends in lockout. Trip operations count every trip, those after failed recloses included."""
    live = _in_2025(reg[reg["function"].isin(["Feeder protection", "Recloser control"])])
    rows = []
    for _, a in live.iterrows():
        rural = a.model == "AT-7"
        shots = int(rng.choice(3, p=[0.0, 0.3, 0.7] if rural else [0.28, 0.47, 0.25]))
        faults = int(rng.poisson(5.1 if rural else 2.6))
        trips = attempts = held = lockouts = 0
        for _ in range(faults):
            trips += 1
            transient = rng.random() < (0.8 if shots else 0.7)
            cleared = False
            for k in range(shots):
                attempts += 1
                if transient and rng.random() < (0.88 if k == 0 else 0.6):
                    held += 1
                    cleared = True
                    break
                trips += 1
            if not cleared:
                lockouts += 1
        rows.append(dict(asset_id=a.asset_id, sub_code=a.sub_code, bay=a.bay, year=2025, reclosing=SHOTS[shots],
                         trip_operations=trips, reclose_attempts=attempts, successful_recloses=held, lockouts=lockouts))
    pd.DataFrame(rows).to_csv(path, index=False)


def test_compliance_xlsx(path, rng, reg, subs):
    """P&C maintenance test compliance for 2025, by model and operating region, for the maintenance programme"""
    cur = dict(zip(subs["sub_code"], subs["region"]))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2025 compliance"
    ws.append(["P&C maintenance test compliance, calendar year 2025"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["P&C maintenance programme report. Devices due a test under P&C Standard 3.2: relays two-yearly at TS "
               "and three-yearly at DS, recloser controls and RTUs every year."])
    hdr = ["Model", "Operating region", "Devices due a test", "Tested on time", "Tested late", "Not tested", "On-time %"]
    ws.append(hdr)
    for j in range(1, len(hdr) + 1):
        ws.cell(row=3, column=j).font = Font(bold=True)
        ws.cell(row=3, column=j).fill = PatternFill("solid", fgColor="E3E9F2")
    live = _in_2025(reg).copy()
    live["op_region"] = [cur[s] for s in live["sub_code"]]
    kind = dict(zip(subs["sub_code"], subs["station_type"]))
    yearly = ("AT-7", "RT-400", "RT-600")
    for (m, r), g in live.groupby(["model", "op_region"]):
        due = int(sum(rng.random() < (1.0 if m in yearly else 0.5 if kind[s] == "TS" else 1 / 3) for s in g["sub_code"]))
        if due == 0:
            continue
        late = int(rng.binomial(due, 0.06))
        nt = int(rng.binomial(due - late, 0.02))
        ok = due - late - nt
        ws.append([m, r, due, ok, late, nt, round(100 * ok / due, 1)])
        ws.cell(row=ws.max_row, column=7).number_format = "0.0"
    for j, w in enumerate([12, 16, 17, 15, 12, 11, 11], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    ws.freeze_panes = "A4"
    wb.save(path)
