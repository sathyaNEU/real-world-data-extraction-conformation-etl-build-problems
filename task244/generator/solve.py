"""task244 golden solution, computed from the shipped pack alone: the published 2027 flows tables read through the
reconciled reinterview's classification error at both quarters (the call), each district's match rate over the year's
four quarter pairs and the reinterview's days from interview to visit (the asks)."""
import datetime as dt
import os

import numpy as np
import pandas as pd
from openpyxl import load_workbook

import estim as E
import params as P

S = P.STATUS
NAME_TO = {"employed": "E", "unemployed": "U", "outside the labour force": "N"}


def code(label):
    return NAME_TO[str(label).strip().lower()]


# ------------------------------------------------------------------ the call
def flows(target, F):
    """the published flows table F1, its district labels filled down"""
    wb = load_workbook(os.path.join(target, F["flows"]), read_only=True, data_only=True)
    ws = wb["F1 flows"]
    rows = list(ws.iter_rows(values_only=True))
    head = next(i for i, r in enumerate(rows) if r and r[0] == "Code")
    cols = list(rows[head])
    out, cur = {}, None
    for r in rows[head + 1:]:
        if r[0] is not None and str(r[0]).startswith("D"):
            cur = str(r[0])
        if cur is None or r[2] is None or str(r[2]).strip().lower() not in NAME_TO:
            if r[0] is not None and not str(r[0]).startswith("D"):
                cur = None
            continue
        a = code(r[2])
        o = out.setdefault(cur, np.zeros((3, 3)))
        for b in S:
            o[S.index(a), S.index(b)] = float(r[cols.index(P.STATUS_NAME[b])])
    return out


def reinterview_table(target, F):
    """reconciled (true) status by the status the original interview recorded, persons reinterviewed in the window"""
    r = pd.read_csv(os.path.join(target, F["reint_results"]))
    t = r["reconciled"] - 1
    o = (r["orig_status"] - 1).astype(int)
    C = np.zeros((3, 3))
    np.add.at(C, (t.values, o.values), 1)
    return C


def call(target, F, C=None, O=None):
    O = O if O is not None else flows(target, F)
    C = C if C is not None else reinterview_table(target, F)
    M = E.misread(C)
    shares = {(d, a + b): E.not_real(O[d], M, a, b) for d in P.DCODES for a, b in E.CELLS}
    return dict(O=O, C=C, M=M, shares=shares, warned=E.warned(shares),
                published={(d, a + b): E.published_rate(O[d], a, b) for d in P.DCODES for a, b in E.CELLS},
                true={(d, a + b): E.true_rate(O[d], M, a, b) for d in P.DCODES for a, b in E.CELLS})


# ------------------------------------------------------------------ the person files
def persons(target, F, dedupe="person"):
    """the five quarterly person files. dedupe="person" (the golden): each person once, a record repeating another of the
    same dwelling and quarter in every field but the household number being the same person received again;
    "exact": identical rows only; "none": as received"""
    fs = [F["persons_" + q.lower()] for q in P.QUARTERS]
    df = pd.concat([pd.read_csv(os.path.join(target, f), dtype={"dwelling": str}) for f in fs], ignore_index=True)
    if dedupe in ("exact", "person"):
        df = df.drop_duplicates(ignore_index=True)
    if dedupe == "person":
        df = df.sort_values(["quarter", "dwelling", "hh", "line"], kind="stable")
        df = df.drop_duplicates([c for c in df.columns if c != "hh"]).sort_index().reset_index(drop=True)
    gz = pd.read_csv(os.path.join(target, F["gazetteer"]))
    moved = gz[gz.district_before.notna() & (gz.district_before.astype(str) != "")]
    old = {(r.district_before, int(r.ea_before)): r.district for r in moved.itertuples()}
    df["d27"] = [old.get((d, int(e)), d) for d, e in zip(df.district, df.ea)]
    df["qi"] = df["quarter"].map({q: i for i, q in enumerate(P.QUARTERS)})
    return df


def match(df, key="both"):
    """each person interviewed at a dwelling due another visit, and whether that person was interviewed again at the
    dwelling's next visit: on the household and line numbers, or, in a household that was not there before, on sex and
    month and year of birth"""
    def has(x, y, cols):
        return pd.MultiIndex.from_frame(x[cols]).isin(pd.MultiIndex.from_frame(y[cols]))
    out = []
    for qi in range(len(P.QUARTERS) - 1):
        a = df[(df.qi == qi) & (df.visit <= 3)].copy()
        b = df[df.qi == qi + 1]
        before = df[df.qi == qi][["dwelling", "hh"]].drop_duplicates()
        bn = b[~has(b, before, ["dwelling", "hh"])]
        birth = ["dwelling", "sex", "birth_year", "birth_month"]
        if key == "line":
            a["matched"] = has(a, b, ["dwelling", "hh", "line"])
        elif key == "birth":
            a["matched"] = has(a, b, birth)
        else:
            a["matched"] = has(a, b, ["dwelling", "hh", "line"]) | has(a, bn, birth)
        out.append(a)
    return pd.concat(out, ignore_index=True)


def ask_match(target, F, df=None, key="both", district="d27"):
    df = df if df is not None else persons(target, F)
    m = match(df, key)
    g = m.groupby(district)["matched"]
    return {d: (int(g.sum()[d]), 100.0 * g.mean()[d], int(g.size()[d])) for d in P.DCODES}


# ------------------------------------------------------------------ the reinterview log
def ask_days(target, F, df=None, date="visit_start", anchor="person", district="d27", late=True,
             quarters=P.QUARTERS[1:]):
    """for the persons of the 2027 selections whose household reinterview was completed (in the window or late), the
    mean days from each person's own interview to the supervisor's completing visit"""
    df = df if df is not None else persons(target, F)
    log = pd.read_csv(os.path.join(target, F["reint_log"]), dtype={"dwelling": str})
    done = log[log.result == "completed"].sort_values("attempt").groupby(["quarter", "dwelling", "hh"]).tail(1)
    if not late:
        r = pd.read_csv(os.path.join(target, F["reint_results"]), dtype={"dwelling": str})
        inw = pd.MultiIndex.from_frame(r[["quarter", "dwelling", "hh"]].drop_duplicates())
        done = done[pd.MultiIndex.from_frame(done[["quarter", "dwelling", "hh"]]).isin(inw)]
    if date == "visit_start":
        done = done.assign(vday=[dt.date.fromisoformat(s[:10]) for s in done.visit_start])
    else:
        done = done.assign(vday=[dt.date.fromisoformat(s) for s in done.attempt_date])
    done = done[done.quarter.isin(quarters)]
    p = df.merge(done[["quarter", "dwelling", "hh", "vday"]], on=["quarter", "dwelling", "hh"])
    iday = p.int_date.map(dt.date.fromisoformat)
    if anchor == "household":
        iday = p.groupby(["quarter", "dwelling", "hh"])["int_date"].transform("min").map(dt.date.fromisoformat)
    p["days"] = [(v - i).days for v, i in zip(p.vday, iday)]
    g = p.groupby(district)["days"]
    return {d: (int(g.size()[d]), g.mean()[d]) for d in P.DCODES}


def answer(target, F):
    A = call(target, F)
    df = persons(target, F)
    A["match"] = ask_match(target, F, df)
    A["days"] = ask_days(target, F, df)
    return A


# ------------------------------------------------------------------ the ladder's routes
def shares_of(fn):
    return {(d, a + b): fn(d, a, b) for d in P.DCODES for a, b in E.CELLS}


def rungs(target, F, A):
    """the routes the ladder and the fork table compare, each a share for every rate"""
    O, C, M = A["O"], A["C"], A["M"]
    r = pd.read_csv(os.path.join(target, F["reint_results"]), dtype={"dwelling": str})
    t = (r.reconciled - 1).values
    o = (r.orig_status - 1).values
    u = (r.reint_status - 1).values
    df = persons(target, F)
    key = ["quarter", "dwelling", "hh", "line"]
    wt = r.merge(df[key + ["wt", "d27", "lfs_status"]].drop_duplicates(key), on=key, how="left")
    Cd = {d: E.table(t[wt.d27.values == d], o[wt.d27.values == d]) for d in P.DCODES}
    q27 = r.quarter.isin(P.QUARTERS[1:]).values
    M27, Mw, Mu = E.misread(E.table(t[q27], o[q27])), E.misread(E.table(t, o, wt.wt.values)), E.misread(E.table(u, o))
    rebuilt = {}
    for name, link in (("bureau key", "line"), ("every person interviewed again", "both")):
        Ob = {d: np.zeros((3, 3)) for d in P.DCODES}
        for qi in range(len(P.QUARTERS) - 1):
            a_ = df[(df.qi == qi) & (df.visit <= 3)]
            b_ = df[df.qi == qi + 1]
            m_ = a_.merge(b_[["dwelling", "hh", "line", "lfs_status"]], on=["dwelling", "hh", "line"], suffixes=("", "_b"))
            if link == "both":
                bn = b_[~pd.MultiIndex.from_frame(b_[["dwelling", "hh"]]).isin(
                    pd.MultiIndex.from_frame(df[df.qi == qi][["dwelling", "hh"]]))]
                rest = a_.merge(m_[["dwelling", "hh", "line"]], how="left", indicator=True)
                rest = rest[rest._merge == "left_only"].drop(columns="_merge")
                m2 = rest.merge(bn[["dwelling", "sex", "birth_year", "birth_month", "lfs_status"]],
                                on=["dwelling", "sex", "birth_year", "birth_month"], suffixes=("", "_b"))
                m_ = pd.concat([m_, m2], ignore_index=True)
            for d in P.DCODES:
                x = m_[m_.d27 == d]
                np.add.at(Ob[d], (x.lfs_status.values - 1, x.lfs_status_b.values - 1), x.wt.values)
        rebuilt[name] = Ob
    out = {
        "R0 predictive values multiplied": shares_of(lambda d, a, b: E.not_real_ppv(C, a, b)),
        "R1 correct readings multiplied": shares_of(lambda d, a, b: E.not_real_both_right(M, a, b)),
        "R1 predictive values by district": shares_of(lambda d, a, b: E.not_real_ppv(Cd[d], a, b)),
        "R2 rate overstatement (net)": shares_of(lambda d, a, b: E.not_real_net(O[d], M, a, b)),
        "R2 flow excess (net)": shares_of(lambda d, a, b: E.not_real_flow_net(O[d], M, a, b)),
        "M 2027Q1 to 2027Q4": shares_of(lambda d, a, b: E.not_real(O[d], M27, a, b)),
        "M weighted": shares_of(lambda d, a, b: E.not_real(O[d], Mw, a, b)),
        "M by district": shares_of(lambda d, a, b: E.not_real(O[d], E.misread(Cd[d]), a, b)),
        "M unreconciled": shares_of(lambda d, a, b: E.not_real(O[d], Mu, a, b)),
        "flows rebuilt, bureau key": shares_of(lambda d, a, b: E.not_real(rebuilt["bureau key"][d], M, a, b)),
        "flows rebuilt, every person": shares_of(
            lambda d, a, b: E.not_real(rebuilt["every person interviewed again"][d], M, a, b)),
    }
    return out, dict(r=r, wt=wt, df=df, Cd=Cd)
