"""task244 independent verifier: recomputes every graded figure from the shipped bundle alone (no generator module is
imported) and checks figures.json, submission.md and the two goldens against it.

    python3 task244/generator/verify.py [target golden figures.json submission.md]
"""
import csv
import datetime as dt
import json
import os
import re
import sys
import zipfile

import numpy as np
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
STATUS = ["employed", "unemployed", "outside the labour force"]
ROWS = {"UE": (1, 0), "EU": (0, 1)}
PERSONS = ["elfs_2026q4_persons.csv", "elfs_2027q1_persons.csv", "elfs_2027q2_persons.csv", "elfs_2027q3_persons.csv",
           "elfs_2027q4_persons.csv"]
QUARTERS = ["2026Q4", "2027Q1", "2027Q2", "2027Q3", "2027Q4"]


def fail(msg):
    print("VERIFY FAIL:", msg)
    sys.exit(1)


def check(cond, msg):
    if not cond:
        fail(msg)


def edge(v, u=0.1):
    return u / 2 - abs(v - round(v / u) * u)


def rows(path):
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def flows(target):
    wb = load_workbook(os.path.join(target, "elfs_2027_annual_flows_tables.xlsx"), read_only=True, data_only=True)
    ws = wb["F1 flows"]
    out, cur, head = {}, None, None
    for r in ws.iter_rows(values_only=True):
        if r and r[0] == "Code":
            head = [str(x).strip().lower() if x else "" for x in r]
            continue
        if head is None:
            continue
        if r[0] and re.fullmatch(r"D\d\d", str(r[0])):
            cur = str(r[0])
        if cur and r[2] and str(r[2]).strip().lower() in STATUS:
            i = STATUS.index(str(r[2]).strip().lower())
            o = out.setdefault(cur, np.zeros((3, 3)))
            for j, s in enumerate(STATUS):
                o[i, j] = float(r[head.index(s)])
        elif r[0] is None and r[2] is None:
            cur = None
    return out


def classification(target):
    C = np.zeros((3, 3))
    for r in rows(os.path.join(target, "elfs_reinterview_results_2026q4_2027q4.csv")):
        C[int(r["reconciled"]) - 1, int(r["orig_status"]) - 1] += 1
    return C


def call(target):
    O, C = flows(target), classification(target)
    M = C / C.sum(1, keepdims=True)
    Mi = np.linalg.inv(M)
    shares, published = {}, {}
    for d, o in O.items():
        Os = o / o.sum()
        Ft = Mi.T @ Os @ Mi
        for k, (i, j) in ROWS.items():
            shares[(d, k)] = 1 - Ft[i, j] * M[i, i] * M[j, j] / Os[i, j]
            published[(d, k)] = o[i, j] / o[i].sum()
    return shares, published, C, M


def persons(target):
    gz = rows(os.path.join(target, "elfs_district_gazetteer_2027.csv"))
    moved = {(g["district_before"], g["ea_before"]): g["district"] for g in gz if g["district_before"]}
    seen, out = set(), []
    for f in PERSONS:
        with open(os.path.join(target, f), newline="", encoding="utf-8") as fh:
            rd = csv.reader(fh)
            head = next(rd)
            ih = head.index("hh")
            recs = sorted(rd, key=lambda x: (x[head.index("quarter")], x[head.index("dwelling")], int(x[ih]),
                                             int(x[head.index("line")])))
            for line in recs:
                t = tuple(v for i, v in enumerate(line) if i != ih)      # each person once
                if t in seen:
                    continue
                seen.add(t)
                r = dict(zip(head, line))
                r["d27"] = moved.get((r["district"], r["ea"]), r["district"])
                out.append(r)
    return out


def match_rates(P):
    byq = {}
    for r in P:
        byq.setdefault(r["quarter"], []).append(r)
    res = {}
    for q, nq in zip(QUARTERS, QUARTERS[1:]):
        nxt = byq[nq]
        hh_before = {(r["dwelling"], r["hh"]) for r in byq[q]}
        lines = {(r["dwelling"], r["hh"], r["line"]) for r in nxt}
        births = {(r["dwelling"], r["sex"], r["birth_year"], r["birth_month"]) for r in nxt
                  if (r["dwelling"], r["hh"]) not in hh_before}
        for r in byq[q]:
            if int(r["visit"]) > 3:
                continue
            hit = (r["dwelling"], r["hh"], r["line"]) in lines or \
                  (r["dwelling"], r["sex"], r["birth_year"], r["birth_month"]) in births
            a = res.setdefault(r["d27"], [0, 0])
            a[0] += hit
            a[1] += 1
    return {d: (v[0], 100.0 * v[0] / v[1], v[1]) for d, v in res.items()}


def days(target, P):
    last = {}
    for r in rows(os.path.join(target, "elfs_reinterview_fieldwork_log_2026q4_2027q4.csv")):
        if r["result"] != "completed" or r["quarter"] == "2026Q4":
            continue
        k = (r["quarter"], r["dwelling"], r["hh"])
        if k not in last or int(r["attempt"]) > last[k][0]:
            last[k] = (int(r["attempt"]), dt.date.fromisoformat(r["visit_start"][:10]))
    res = {}
    for r in P:
        k = (r["quarter"], r["dwelling"], r["hh"])
        if k in last:
            a = res.setdefault(r["d27"], [0, 0])
            a[0] += (last[k][1] - dt.date.fromisoformat(r["int_date"])).days
            a[1] += 1
    return {d: (v[1], v[0] / v[1]) for d, v in res.items()}


def docx_text(path):
    with zipfile.ZipFile(path) as z:
        x = z.read("word/document.xml").decode("utf-8")
    return re.sub(r"<[^>]+>", "", x)


def main(target, golden, figpath, subpath):
    fig = json.load(open(figpath))
    sub = open(subpath, encoding="utf-8").read()
    shares, published, C, M = call(target)
    warned = sorted(k for k, v in shares.items() if v >= 1 / 3)
    n = 0
    for c in fig["cells"]:
        k = (c["district"], c["key"])
        check(abs(round(100 * shares[k], 1) - c["not_real"]) < 1e-9, "share %s %s" % k)
        check(abs(round(100 * published[k], 1) - c["published"]) < 1e-9, "published %s %s" % k)
        check(c["warning"] == (k in warned), "warning %s %s" % k)
        check(edge(100 * shares[k]) >= 0.007, "share %s %s near its edge" % k)
        line = "%s (%s) %s: %.1f, %s" % (c["name"], c["district"], c["flow"], c["not_real"],
                                         "warning" if c["warning"] else "no warning")
        check(line in sub, "submission line: " + line)
        n += 5
    check(fig["call"]["count"] == len(warned) and fig["call"]["of"] == 24, "the call's count")
    check(int(C.sum()) == fig["reinterview"]["persons"], "persons reinterviewed")
    st = ["E", "U", "N"]
    for k, v in fig["reinterview"]["misread"].items():
        check(abs(round(100 * M[st.index(k[0]), st.index(k[1])], 2) - v) < 1e-9, "misreading %s" % k)
        check(edge(100 * M[st.index(k[0]), st.index(k[1])], 0.01) >= 0.0007, "misreading %s near its edge" % k)
        n += 2
    ppv = np.diag(C) / C.sum(0)
    check(abs(round(100 * (1 - ppv[0] * ppv[1]), 1) - fig["rungs"]["ppv_statewide"]) < 1e-9, "the default path's figure")
    check(abs(round(100 * (1 - M[0, 0] * M[1, 1]), 1) - fig["rungs"]["both_right"]) < 1e-9, "the correct readings' figure")
    n += 3
    P = persons(target)
    mr = match_rates(P)
    for x in fig["match"]:
        g = mr[x["district"]]
        check(g[0] == x["persons"] and abs(round(g[1], 1) - x["rate"]) < 1e-9 and g[2] == x["due"],
              "match %s" % x["district"])
        check(edge(g[1]) >= 0.007, "match %s near its edge" % x["district"])
        check("   - %s: %.1f, %s persons" % (x["name"], x["rate"], "{:,}".format(x["due"])) in sub,
              "submission match line %s" % x["district"])
        n += 3
    dd = days(target, P)
    for x in fig["days"]:
        g = dd[x["district"]]
        check(g[0] == x["persons"] and abs(round(g[1], 1) - x["mean"]) < 1e-9, "days %s" % x["district"])
        check(edge(g[1]) >= 0.007, "days %s near its edge" % x["district"])
        check("   - %s: %.1f days" % (x["name"], x["mean"]) in sub, "submission days line %s" % x["district"])
        n += 3
    # the goldens carry the same figures
    wb = load_workbook(os.path.join(golden, "elfs_2027_flows_quality.xlsx"), read_only=True, data_only=True)
    got = {(r[0], r[2]): r for r in wb["Warning check"].iter_rows(min_row=5, values_only=True) if r[0] and
           re.fullmatch(r"D\d\d", str(r[0]))}
    check(len(got) == 24, "the golden workbook's 24 rows")
    for c in fig["cells"]:
        r = got[(c["district"], c["flow"].capitalize() + " rate")]
        check(abs(r[5] - c["not_real"]) < 1e-9 and (r[6] == "Warning") == c["warning"], "golden row %s %s" %
              (c["district"], c["key"]))
        n += 1
    m = {r[0]: r for r in wb["Match rate"].iter_rows(min_row=4, values_only=True) if r[0] and re.fullmatch(r"D\d\d", str(r[0]))}
    t = {r[0]: r for r in wb["Reinterview timing"].iter_rows(min_row=4, values_only=True) if r[0] and
         re.fullmatch(r"D\d\d", str(r[0]))}
    for x in fig["match"]:
        check(m[x["district"]][2] == x["due"] and m[x["district"]][3] == x["persons"] and
              abs(m[x["district"]][4] - x["rate"]) < 1e-9, "golden match")
    for x in fig["days"]:
        check(abs(t[x["district"]][3] - x["mean"]) < 1e-9, "golden days")
    n += 24
    note = docx_text(os.path.join(golden, "elfs_2027_flows_note.docx"))
    for c in fig["cells"]:
        if c["warning"]:
            check(c["name"] in note, "the note names %s" % c["name"])
    check(("%d of the %d" % (len(warned), 24)) in note, "the note gives the count")
    check("—" not in sub + note and "–" not in sub + note, "no dashes")
    n += 3
    print("verify: %d checks passed" % n)
    return n


if __name__ == "__main__":
    a = sys.argv[1:]
    if len(a) == 4:
        main(*a)
    else:
        main(os.path.join(TASK, "target"), os.path.join(TASK, "golden"), os.path.join(HERE, "figures.json"),
             os.path.join(TASK, "submission.md"))
