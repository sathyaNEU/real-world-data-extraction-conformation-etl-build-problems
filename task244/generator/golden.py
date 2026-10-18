"""task244 goldens: the quality workbook (each of the 24 district rates' share of counted moves that are not real moves
and its warning, the reinterview's classification error, each district's match rate and reinterview timing) and the
quality note to the management meeting with its chart."""
import io
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import params as P
import postprocess

XLSX = "elfs_2027_flows_quality.xlsx"
DOCX = "elfs_2027_flows_note.docx"
ORG = "%s, %s" % (P.BUREAU, P.DIVISION)
AUTHOR = "Comfort Olawale"
FLOW_COL = {"UE": "#2F6B8F", "EU": "#B5532B"}
FLOW_LIGHT = {"UE": "#C9DCE8", "EU": "#EFCDBE"}


def _head(ws, row, head, height=46):
    thin = Side(style="thin", color="A6ADB8")
    for j, h in enumerate(head, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(bold=True, size=9.5)
        c.fill = PatternFill("solid", fgColor="DDE4EE")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = Border(bottom=thin)
    ws.row_dimensions[row].height = height


def _notes(ws, row, notes, ncol, heights):
    for k, t in enumerate(notes):
        c = ws.cell(row=row + k, column=1, value=t)
        c.font = Font(size=8.5, italic=(k == len(notes) - 1))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row + k, start_column=1, end_row=row + k, end_column=ncol)
        ws.row_dimensions[row + k].height = heights[k]


def _widths(ws, widths):
    for col, w in zip("ABCDEFGHIJ", widths):
        ws.column_dimensions[col].width = w


def _workbook(path, F):
    wb = Workbook()
    ws = wb.active
    ws.title = "Warning check"
    ws["A1"] = "2027 annual labour market flows: classification-error check of the 24 district rates"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("%s, persons aged 15 and over, survey districts as from 1 April 2027, average of the four quarter pairs "
                "2026Q4 to 2027Q4. Warning as defined in the methodology note, section 8.1." % P.SURVEY)
    ws["A2"].font = Font(italic=True, size=9)
    head = ["Code", "District", "Rate", "Published rate (%)", "Rate with the classification error taken out (%)",
            "Share of counted moves that are not real moves (%)", "Classification-error warning"]
    _head(ws, 4, head, 58)
    for i, c in enumerate(F["cells"], 5):
        vals = [c["district"], c["name"], c["flow"].capitalize() + " rate", c["published"], c["true_rate"],
                c["not_real"], "Warning" if c["warning"] else "No warning"]
        for j, v in enumerate(vals, 1):
            x = ws.cell(row=i, column=j, value=v)
            x.font = Font(size=10, bold=(j == 7 and c["warning"]))
            if j in (4, 5, 6):
                x.number_format = "0.0"
            x.alignment = Alignment(horizontal="left" if j in (1, 2, 3) else ("center" if j == 7 else "right"))
            if c["warning"]:
                x.fill = PatternFill("solid", fgColor="FBEFD9")
    n = 5 + len(F["cells"])
    _notes(ws, n + 1, [
        "Published rate: Table F2 of the 2027 annual flows tables, the persons Table F1 counts as moving over the persons "
        "in the first status in the first quarter. Rate with the classification error taken out: the same rate on the "
        "true flows, the Table F1 flows with the reinterview's classification error (sheet Classification error) taken "
        "out at both quarters of each pair.",
        "Share of counted moves that are not real moves: of the persons Table F1 counts as moving from the first status "
        "to the second, the share whose reconciled statuses in the two quarters were not that pair; one less the true "
        "flow recorded correctly at both quarters, over the counted flow. Warning: a third or more (section 8.1).",
        "Sources: 2027 annual flows tables as tabulated; reinterview results, 2026Q4 to 2027Q4; reinterview protocol, "
        "version 3."], len(head), [44, 44, 18])
    _widths(ws, [7, 21, 18, 11, 15, 15, 14])
    ws.freeze_panes = "C5"
    ws.print_title_rows = "4:4"

    ce = wb.create_sheet("Classification error")
    ce["A1"] = "Reinterview, interviews of 2026Q4 to 2027Q4: reconciled status by the status the interview recorded"
    ce["A1"].font = Font(bold=True, size=11)
    names = [P.STATUS_NAME[s] for s in P.STATUS]
    _head(ce, 3, ["Reconciled status", "Recorded " + names[0].lower(), "Recorded " + names[1].lower(),
                  "Recorded " + names[2].lower(), "Persons"], 34)
    m = F["reinterview"]
    for i, s in enumerate(P.STATUS, 4):
        row = m["table"][s]
        for j, v in enumerate([P.STATUS_NAME[s]] + [row[t] for t in P.STATUS] + [sum(row.values())], 1):
            x = ce.cell(row=i, column=j, value=v)
            if j > 1:
                x.number_format = "#,##0"
    tot = ce.cell(row=7, column=1, value="All persons reinterviewed")
    tot.font = Font(bold=True)
    x = ce.cell(row=7, column=5, value=m["persons"])
    x.number_format = "#,##0"
    x.font = Font(bold=True)
    _head(ce, 9, ["Reconciled status", "Recorded " + names[0].lower() + " (%)", "Recorded " + names[1].lower() + " (%)",
                  "Recorded " + names[2].lower() + " (%)", ""], 34)
    for i, s in enumerate(P.STATUS, 10):
        vals = [P.STATUS_NAME[s]] + [m["matrix"][s + t] for t in P.STATUS]
        for j, v in enumerate(vals, 1):
            x = ce.cell(row=i, column=j, value=v)
            if j > 1:
                x.number_format = "0.00"
                x.font = Font(bold=(s == P.STATUS[j - 2]))
    _notes(ce, 14, [
        "Persons reinterviewed within the 14-day window, each counted once and unweighted, for the state over the five "
        "quarters whose interviews enter the 2027 tables (reinterview protocol, section 7). Each row gives the chance a "
        "person of that reconciled status is recorded as each status; the diagonal is the chance of a correct record. "
        "The shares carry these chances unrounded.",
        "Source: reinterview results, 2026Q4 to 2027Q4."], 5, [44, 16])
    _widths(ce, [26, 16, 16, 18, 12])

    mr = wb.create_sheet("Match rate")
    mr["A1"] = "Match rate by district, the four quarter pairs 2026Q4 to 2027Q4 (section 8.2)"
    mr["A1"].font = Font(bold=True, size=11)
    _head(mr, 3, ["Code", "District", "Persons interviewed in a dwelling due another visit, four pairs summed",
                  "Persons interviewed again at the next visit, four pairs summed", "Match rate (%)"], 58)
    for i, x in enumerate(F["match"], 4):
        for j, v in enumerate([x["district"], x["name"], x["due"], x["persons"], x["rate"]], 1):
            c = mr.cell(row=i, column=j, value=v)
            c.number_format = "0.0" if j == 5 else ("#,##0" if j in (3, 4) else "General")
    _notes(mr, 4 + len(F["match"]) + 1, [
        "A person counts as interviewed again when found at the dwelling's next visit in the same household on the same "
        "line, or in a household new to the dwelling with the same sex and month and year of birth (households that "
        "split renumber their lines). Districts as from 1 April 2027 (district gazetteer). Each person once: the "
        "records of 14 May 2027 that the server took in again on 17 May under new household numbers, and the "
        "second-quarter batch of Market Towns area 5 received twice, counted once.",
        "Source: the quarterly person files, 2026Q4 to 2027Q4."], 5, [44, 16])
    _widths(mr, [7, 21, 18, 18, 11])

    rt = wb.create_sheet("Reinterview timing")
    rt["A1"] = "Days from interview to reinterview by district, interviews of 2027Q1 to 2027Q4 (section 8.3)"
    rt["A1"].font = Font(bold=True, size=11)
    _head(rt, 3, ["Code", "District", "Persons whose household reinterview was completed",
                  "Mean days from interview to reinterview"], 46)
    for i, x in enumerate(F["days"], 4):
        for j, v in enumerate([x["district"], x["name"], x["persons"], x["mean"]], 1):
            c = rt.cell(row=i, column=j, value=v)
            c.number_format = "0.0" if j == 4 else ("#,##0" if j == 3 else "General")
    _notes(rt, 4 + len(F["days"]) + 1, [
        "Each person's own interview date to the date of the supervisor's visit that completed the household's "
        "reinterview, within the window or after it. The visit's date is the date the reinterview form was opened "
        "(visit_start): from the tablet app update, which each supervisor's tablet took at a Friday sync between 18 "
        "June and 9 July 2027, the log's call date is the date the tablet next synchronised, the Friday after the "
        "call.",
        "Source: reinterview fieldwork log, 2026Q4 to 2027Q4; the quarterly person files."], 4, [56, 16])
    _widths(rt, [7, 21, 20, 18])

    ns = wb.create_sheet("Notes")
    lines = [
        ("As of", "10 February 2028; the 2027 annual flows tables as tabulated on 28 January 2028"),
        ("Warning", "A rate carries the classification-error warning when a third or more of the moves it counts are "
                    "not real moves, a person's status in a quarter being the reinterview's reconciled status "
                    "(methodology note, section 8.1)."),
        ("Rates", "The 24 rates of the annual tables: each district's job-finding rate (unemployed to employed) and "
                  "job-loss rate (employed to unemployed), the average of the four quarter pairs 2026Q4 to 2027Q4."),
        ("Classification error", "Persons reinterviewed within the window over the interviews of 2026Q4 to 2027Q4, "
                                 "each once and unweighted, for the state (reinterview protocol, section 7)."),
        ("Match rate", "Section 8.2. Persons in a dwelling's fourth visit are not due another visit and are left out."),
        ("Reinterview timing", "Section 8.3. Reinterviews completed after the window are in; selected households "
                               "never reached, refused or not contacted are out."),
        ("Districts", "As from 1 April 2027 for every quarter (methodology note, section 6)."),
    ]
    for i, (k, v) in enumerate(lines, 1):
        ns.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = ns.cell(row=i, column=2, value=v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ns.row_dimensions[i].height = 30 if len(v) > 100 else 16
    ns.column_dimensions["A"].width = 20
    ns.column_dimensions["B"].width = 100
    for sh in wb.worksheets:
        sh.page_setup.orientation = "landscape"
        sh.page_setup.paperSize = sh.PAPERSIZE_A4
        sh.sheet_properties.pageSetUpPr.fitToPage = True
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
    wb.save(path)
    postprocess.scrub(path, AUTHOR, "2028-02-10")
    postprocess.finish_office(path, "2028-02-03T08:20:00Z", "2028-02-10T14:05:00Z", ORG)


def chart_png(F):
    cells = {(c["district"], c["key"]): c for c in F["cells"]}
    names = [x["name"] for x in F["match"]]
    codes = [x["district"] for x in F["match"]]
    fig, ax = plt.subplots(figsize=(6.8, 5.0), dpi=200)
    y0 = list(range(len(codes)))[::-1]
    h = 0.38
    for yy, d in zip(y0, codes):
        for k, off in (("UE", h / 2), ("EU", -h / 2)):
            c = cells[(d, k)]
            col = FLOW_COL[k] if c["warning"] else FLOW_LIGHT[k]
            ax.barh(yy + off, c["not_real"], height=h, color=col, edgecolor=FLOW_COL[k], lw=0.6, zorder=2)
            ax.text(c["not_real"] + 0.8, yy + off, "%.1f" % c["not_real"], va="center", fontsize=6.2,
                    color="#222222" if c["warning"] else "#666666", fontweight="bold" if c["warning"] else "normal")
    line = F["constants"]["warn_share"]
    ax.axvline(100.0 / 3.0, color="#1F2A3A", lw=1.1, ls="--", zorder=3)
    ax.text(100.0 / 3.0 + 0.8, len(codes) - 0.35, "Warning: a third of counted moves (%.1f%%)" % line, fontsize=7,
            color="#1F2A3A", va="center")
    ax.set_yticks(y0)
    ax.set_yticklabels(names, fontsize=7.8)
    ax.set_xlim(0, 100)
    ax.set_ylim(-0.7, len(codes) - 0.05)
    ax.set_xlabel("Share of the rate's counted moves that are not real moves, per cent", fontsize=7.8)
    ax.set_title("%d of the %d district rates carry the classification-error warning" %
                 (F["call"]["count"], F["constants"]["rates"]), fontsize=9.5, loc="left", fontweight="bold")
    from matplotlib.patches import Patch
    ax.legend(handles=[Patch(color=FLOW_COL["UE"], label="Job-finding rate, warning"),
                       Patch(facecolor=FLOW_LIGHT["UE"], edgecolor=FLOW_COL["UE"], label="Job-finding rate, no warning"),
                       Patch(color=FLOW_COL["EU"], label="Job-loss rate, warning"),
                       Patch(facecolor=FLOW_LIGHT["EU"], edgecolor=FLOW_COL["EU"], label="Job-loss rate, no warning")],
              fontsize=6.4, loc="lower right", frameon=False)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.tick_params(axis="x", labelsize=7.5)
    fig.text(0.01, 0.012, "The published flows read through the reconciled reinterview's classification error at both "
                          "quarters of each pair (methodology note, section 8.1).\nSource: %s, 2027 annual flows tables "
                          "and reinterview results, 2026Q4 to 2027Q4." % P.SURVEY_SHORT, fontsize=5.8, color="#555555",
             va="bottom")
    fig.tight_layout(rect=(0, 0.06, 1, 1))
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    return buf.getvalue()


def _note(path, F, G):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    hdr = sec.header.paragraphs[0]
    hdr.text = ORG
    hdr.runs[0].font.size = Pt(8)
    ftr = sec.footer.paragraphs[0]
    ftr.text = "For the %s of 16 February 2028    Page " % P.FORUM
    ftr.runs[0].font.size = Pt(8)
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = ftr.add_run()
    run.font.size = Pt(8)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)

    def para(text, bold=False, size=None, after=6, italic=False, color=None):
        p = d.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        p.paragraph_format.space_after = Pt(after)
        return p

    para("QUALITY NOTE: 2027 ANNUAL LABOUR MARKET FLOWS", bold=True, size=9, after=2, color="555555")
    para("To: the %s, through Atinuke Adebisi, Director    From: Comfort Olawale, survey methodologist    Date: 10 "
         "February 2028" % P.FORUM, size=9, after=10, color="555555")
    para(G["lead"], bold=True, size=13, after=8)
    for t in G["paras"][:1]:
        para(t, after=8)
    d.add_picture(io.BytesIO(chart_png(F)), width=Inches(6.4))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    warned = [c for c in F["cells"] if c["warning"]]
    tb = d.add_table(rows=1, cols=5)
    tb.style = "Light Grid Accent 1"
    for j, h in enumerate(["District", "Rate", "Published (%)", "Error taken out (%)", "Counted moves not real (%)"]):
        tb.rows[0].cells[j].text = h
    for c in warned:
        cells = tb.add_row().cells
        for j, v in enumerate([c["name"], c["flow"].capitalize(), "%.1f" % c["published"], "%.1f" % c["true_rate"],
                               "%.1f" % c["not_real"]]):
            cells[j].text = v
    for row in tb.rows:
        for j, cell in enumerate(row.cells):
            for p_ in cell.paragraphs:
                p_.paragraph_format.space_after = Pt(0)
                if j >= 2:
                    p_.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for r_ in p_.runs:
                    r_.font.size = Pt(8.5)
    para("The %d rates that carry the warning. Published: Table F2. Error taken out: the same rate on the true flows."
         % len(warned), italic=True, size=8, after=10, color="555555")
    for t in G["paras"][1:]:
        para(t, after=8)
    para(G["footnote"], size=8.5, after=6, color="444444")
    para(G["sources"], italic=True, size=8.5, after=0)
    d.core_properties.author = AUTHOR
    d.core_properties.title = G["title"]
    d.core_properties.comments = ""
    d.save(path)
    postprocess.scrub(path, AUTHOR, "2028-02-10")
    postprocess.finish_office(path, "2028-02-08T09:15:00Z", "2028-02-10T15:30:00Z", ORG)


def write(folder, F, G):
    os.makedirs(folder, exist_ok=True)
    _workbook(os.path.join(folder, XLSX), F)
    _note(os.path.join(folder, DOCX), F, G)
