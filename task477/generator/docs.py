"""The pack's documents: the fund's guidelines and the board's performance reporting manual (PDF), the state's county
file layout and notes (PDF), the area and industry titles (CSV), the training provider list (xlsx), and two
distractors (the industry projections and the occupational wage estimates, xlsx)."""
from reportlab import rl_config
rl_config.invariant = 1
from reportlab.lib import colors                                    # noqa: E402
from reportlab.lib.pagesizes import letter                          # noqa: E402
from reportlab.lib.styles import ParagraphStyle                     # noqa: E402
from reportlab.lib.units import inch                                # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle  # noqa: E402
from openpyxl import Workbook                                       # noqa: E402
from openpyxl.styles import Font                                    # noqa: E402

import numpy as np                                                  # noqa: E402

import world as W   # noqa: E402

SECTIONS = {"eligible": "2.1", "measure": "2.2", "slate": "2.3", "exiters": "3.2", "ita": "3.3"}
BODY = ParagraphStyle("b", fontName="Times-Roman", fontSize=10.5, leading=13.5, spaceAfter=6)
HEAD = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=11, leading=14, spaceBefore=8, spaceAfter=4)
TITLE = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=15, leading=19, spaceAfter=4)
SUB = ParagraphStyle("s", fontName="Helvetica", fontSize=9, leading=12, textColor=colors.HexColor("#444444"),
                     spaceAfter=10)
SMALL = ParagraphStyle("sm", fontName="Times-Roman", fontSize=9.5, leading=12, spaceAfter=4)


def _doc(path, title, org):
    d = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.95 * inch, rightMargin=0.95 * inch,
                          topMargin=0.9 * inch, bottomMargin=0.9 * inch, title=title, author=org,
                          creator=org, producer=org, subject=title)

    def frame(c, d_):
        c.saveState()
        c.setFont("Helvetica", 8)
        c.setFillColor(colors.HexColor("#555555"))
        c.drawString(0.95 * inch, 0.55 * inch, "%s  |  %s" % (org, title))
        c.drawRightString(letter[0] - 0.95 * inch, 0.55 * inch, "Page %d" % d_.page)
        c.restoreState()
    return d, frame


def _table(rows, widths, size=9):
    t = Table(rows, colWidths=widths)
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, 0), "Helvetica-Bold", size), ("FONT", (0, 1), (-1, -1), "Times-Roman", size + .5),
                           ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.black), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


def write_guidelines(path):
    org = "State Workforce Development Council"
    title = "%s: 2027-2029 Round Guidelines" % W.FUND
    d, frame = _doc(path, title, org)
    s = [Paragraph("%s<br/>2027-2029 Round: Guidelines for Local Workforce Boards" % W.FUND, TITLE),
         Paragraph("%s. Issued %s." % (org, W.GUIDELINE_DATE), SUB),
         Paragraph("1. Purpose", HEAD),
         Paragraph("The Fund supports industry-led partnerships in which a local workforce development board, employers "
                   "and training providers in one industry plan and deliver training for that industry's jobs. Each "
                   "board names a slate of industries for the 2027-2029 round; the Fund supports up to four "
                   "partnerships per board for three program years from 1 July 2027.", BODY),
         Paragraph("2. The industry slate", HEAD),
         Paragraph("2.1 <b>Eligible industries.</b> An industry is a private-ownership NAICS sector as the state's "
                   "Quarterly Census of Employment and Wages (QCEW) program publishes it for counties (sectors 31-33, "
                   "44-45 and 48-49 each count as one sector). An industry is eligible if its annual average jobs in "
                   "the board's area in 2025 were 1,000 or more. Unclassified establishments (NAICS 99) are not an "
                   "industry for this purpose.", BODY),
         Paragraph("2.2 <b>Growth in pay per job.</b> The Fund ranks eligible industries on their growth in pay per "
                   "job from 2021, the Fund's base year, to 2025. An industry's pay per job for a year is the pay its "
                   "employees earned for the work they performed in that calendar year, divided by its annual average "
                   "jobs that year (the mean of the twelve monthly job counts). Growth is the 2025 pay per job "
                   "divided by the 2021 pay per job, less one.", BODY),
         Paragraph("2.3 <b>The slate.</b> A board's slate is its four eligible industries with the highest growth in "
                   "pay per job. A board with fewer than four eligible industries names them all.", BODY),
         Paragraph("2.4 <b>Source.</b> Boards use the state's published QCEW county data for their area. Boards whose "
                   "area is several counties combine the counties' published figures.", BODY),
         Paragraph("3. What the application reports", HEAD),
         Paragraph("3.1 The slate, in order of growth, and the eligible industry next in line.", BODY),
         Paragraph("3.2 For every eligible industry, the number of the board's program exiters in calendar 2025 who "
                   "were employed in that industry in the second quarter after exit, counted as the board's "
                   "performance reporting manual counts employment after exit.", BODY),
         Paragraph("3.3 For every eligible industry, the board's individual training account (ITA) payments in "
                   "calendar 2025 for training aimed at that industry. Training is aimed at the industry the state's "
                   "eligible training provider list gives as the program's target industry. Payments are the checks "
                   "the board's fiscal agent issued to providers in the calendar year, less any of those checks later "
                   "voided.", BODY),
         Paragraph("4. Timetable", HEAD),
         Paragraph("Boards vote on their slate by 30 November 2026 and file the application with the Council by "
                   "15 January 2027. The Council announces awards in April 2027.", BODY)]
    d.build(s, onFirstPage=frame, onLaterPages=frame)


def write_layout(path, n):
    org = W.AGENCY
    title = "QCEW County Data Files: Layout and Notes"
    d, frame = _doc(path, title, org)
    rows = [["Field", "Description"],
            ["area_fips", "County FIPS code (state 52)."],
            ["own_code", "Ownership: 0 total covered, 1 federal government, 2 state government, 3 local government, "
                         "5 private."],
            ["industry_code", "NAICS code; 10 is the total of all industries."],
            ["agglvl_code", "70 county total; 71 county by ownership; 74 county, ownership and NAICS sector; "
                            "75 county, ownership and NAICS 3-digit subsector."],
            ["size_code", "0 (all establishment sizes)."],
            ["year, qtr", "Reference year and calendar quarter."],
            ["disclosure_code", "N where the cell is not disclosed to protect an employer's confidentiality; the "
                                "cell's values are then zero."],
            ["qtrly_estabs", "Establishments reporting in the quarter (qtrly_estabs_count before 2020)."],
            ["month1_emplvl to month3_emplvl", "Covered workers who worked during, or received pay for, the pay "
                                               "period including the 12th of each month of the quarter."],
            ["total_qtrly_wages", "Total wages paid to covered workers during the calendar quarter, regardless of "
                                  "when the services were performed. Includes bonuses, stock options, severance pay, "
                                  "the cash value of meals and lodging, and tips and other gratuities."],
            ["taxable_qtrly_wages", "Wages subject to the state unemployment insurance tax (the first $15,000 per "
                                    "worker in a calendar year)."],
            ["qtrly_contributions", "Unemployment insurance contributions due on the quarter's taxable wages."],
            ["avg_wkly_wage", "Total quarterly wages divided by the average of the three monthly employment levels, "
                              "divided by 13, rounded to the dollar."],
            ["oty_month1_emplvl_chg, oty_total_qtrly_wages_chg", "Change from the same quarter a year earlier "
                                                                  "(from 2020; zero where either quarter is not "
                                                                  "disclosed)."]]
    s = [Paragraph("Quarterly Census of Employment and Wages<br/>County Data Files: Layout and Notes", TITLE),
         Paragraph("%s. Revised June 2026." % org, SUB),
         Paragraph("The county data files carry every county of the state by ownership, sector and 3-digit "
                   "subsector, one row per cell per quarter. Data come from the quarterly contribution reports "
                   "employers subject to the state's unemployment insurance law file with the Department.", BODY),
         _table([[Paragraph(a, SMALL), Paragraph(b, SMALL)] for a, b in rows[1:]][:0] or
                [[a, Paragraph(b, SMALL)] if i else [a, b] for i, (a, b) in enumerate(rows)],
                [2.0 * inch, 4.6 * inch]),
         Paragraph("Files and releases", HEAD),
         Paragraph("%s holds 2013 through 2019 in the layout used before 2020. %s holds the first quarter of 2020 "
                   "through the third quarter of 2025 as published by the March 2026 release. %s is the 2025 annual "
                   "release (June 2026): it carries all four quarters of 2025 and replaces the figures first "
                   "published for the first three quarters, which are revised for late and amended reports."
                   % (W.FILES["qcew_a"], W.FILES["qcew_b"], W.FILES["qcew_c"]), BODY),
         Paragraph("Industry coding", HEAD),
         Paragraph("Data from the first quarter of 2022 are coded to NAICS 2022, earlier data to NAICS 2017 or NAICS "
                   "2012. Under NAICS 2022, retail subsectors 452 and 454 were replaced by 455 to 459; sector totals "
                   "are not affected.", BODY),
         Paragraph("Disclosure", HEAD),
         Paragraph("A suppressed subsector is still included in its sector, ownership and county totals, so a "
                   "sector's figures can exceed the sum of its disclosed subsectors.", BODY)]
    d.build(s, onFirstPage=frame, onLaterPages=frame)


def write_manual(path):
    org = W.BOARD
    title = "Performance Reporting Manual"
    d, frame = _doc(path, title, org)
    s = [Paragraph("%s<br/>Performance Reporting Manual (excerpt)" % W.BOARD, TITLE),
         Paragraph("Section 4, Employment after exit. Revised February 2025.", SUB),
         Paragraph("4.1 <b>Exiters.</b> A participant exits on the last date of a service funded by a WIOA Title I "
                   "program after which no further service is planned. Participants enrolled in more than one Title I "
                   "program (co-enrolled) exit each program on the same date and are one exiter. A participant who "
                   "exits more than once in a calendar year is counted once for that year, at the last exit.", BODY),
         Paragraph("4.2 <b>Employment in the second quarter after exit.</b> An exiter is employed in the second "
                   "quarter after the exit quarter (an exit in January to March is followed by July to September) "
                   "when the state's wage-record match return shows wages above zero paid to the exiter in that "
                   "quarter.", BODY),
         Paragraph("4.3 <b>Industry of employment.</b> An exiter employed in the quarter is counted in the industry "
                   "of the business where the exiter worked, by the NAICS code the state's employer accounts file "
                   "gives that business. An exiter paid by more than one business in the quarter is counted once, in "
                   "the industry of the business that paid the exiter the most.", BODY),
         Paragraph("4.4 <b>Sources.</b> Exits come from the case management system's exit extract; wages from the "
                   "state's quarterly wage-record match return; employer industry from the state's employer accounts "
                   "extract.", BODY)]
    d.build(s, onFirstPage=frame, onLaterPages=frame)


def area_rows():
    rows = [["area_fips", "area_title"], ["52000", "Statewide"]]
    allc = [(f, n) for n, f, _ in W.COUNTIES] + [(W.COUNTY_FIPS, W.COUNTY)]
    for f, n in sorted(allc):
        rows.append([f, n])
    return rows


def industry_rows(x):
    import records as R
    codes = {}
    for sec, (title, *_rest) in W.SECTORS.items():
        codes[sec] = title
    titles3 = {"111": "Crop production", "115": "Agriculture and forestry support activities", "212": "Mining (except oil and gas)",
               "221": "Utilities", "236": "Construction of buildings", "237": "Heavy and civil engineering construction",
               "238": "Specialty trade contractors", "311": "Food manufacturing", "326": "Plastics and rubber products manufacturing",
               "332": "Fabricated metal product manufacturing", "333": "Machinery manufacturing",
               "336": "Transportation equipment manufacturing", "423": "Merchant wholesalers, durable goods",
               "424": "Merchant wholesalers, nondurable goods", "425": "Wholesale trade agents and brokers",
               "441": "Motor vehicle and parts dealers", "444": "Building material and garden supply stores",
               "445": "Food and beverage retailers", "447": "Gasoline stations", "452": "General merchandise stores",
               "454": "Nonstore retailers", "455": "General merchandise retailers",
               "456": "Health and personal care retailers", "458": "Clothing, clothing accessories, shoe, and jewelry retailers",
               "459": "Sporting goods, hobby, musical instrument, book, and miscellaneous retailers",
               "484": "Truck transportation", "492": "Couriers and messengers", "493": "Warehousing and storage",
               "517": "Telecommunications", "518": "Computing infrastructure providers, data processing, web hosting",
               "519": "Web search portals, libraries, archives, and other information services",
               "522": "Credit intermediation and related activities", "523": "Securities, commodity contracts, investments",
               "524": "Insurance carriers and related activities", "531": "Real estate", "532": "Rental and leasing services",
               "541": "Professional, scientific, and technical services", "551": "Management of companies and enterprises",
               "561": "Administrative and support services", "562": "Waste management and remediation services",
               "611": "Educational services", "621": "Ambulatory health care services", "622": "Hospitals",
               "623": "Nursing and residential care facilities", "624": "Social assistance",
               "711": "Performing arts and spectator sports", "713": "Amusement, gambling, and recreation industries",
               "721": "Accommodation", "722": "Food services and drinking places", "811": "Repair and maintenance",
               "812": "Personal and laundry services", "813": "Membership associations and organizations",
               "999": "Unclassified"}
    rows = [["industry_code", "industry_title"], ["10", "Total, all industries"]]
    for c in sorted(set(codes) | set(titles3), key=lambda c: (c.split("-")[0].zfill(3), c)):
        rows.append([c, ("NAICS %s %s" % (c, codes[c])) if c in codes else "NAICS %s %s" % (c, titles3[c])])
    return rows


def write_etpl(path, x):
    import records as R
    wb = Workbook()
    ws = wb.active
    ws.title = "Programs"
    ws.append(["Eligible Training Provider List: programs approved for Tallis Valley ITAs"])
    ws.append(["Extract of 30 June 2026. Target industry is the industry the provider's application names."])
    ws.append([])
    ws.append(["Program ID", "Provider", "Program", "CIP", "Target Industry (NAICS)", "Approved", "Status"])
    ws["A4"].font = Font(bold=True)
    for p in sorted(x["asks"]["programs"], key=lambda p: p["program_id"]):
        ws.append([p["program_id"], p["provider"], p["name"], p["cip"], "%s %s" % (p["sec"], W.SECTORS[p["sec"]][0]),
                   p["approved"].strftime("%Y-%m-%d"), "Active"])
    for col, w in zip("ABCDEFG", (11, 34, 44, 9, 62, 12, 8)):
        ws.column_dimensions[col].width = w
    wb.save(path)


def write_projections(path, x):
    rng = np.random.default_rng(W.SEED + 7)
    wb = Workbook()
    ws = wb.active
    ws.title = "Industry"
    ws.append(["Tallis County industry employment projections, 2024-2034"])
    ws.append(["%s. Long-term projections, published May 2026." % W.AGENCY])
    ws.append([])
    ws.append(["NAICS", "Industry", "Employment 2024", "Projected 2034", "Change", "Percent change"])
    for sec, (title, jobs, *_r) in W.SECTORS.items():
        if sec == "99":
            continue
        e = int(round(jobs * float(np.exp(rng.normal(-.01, .03))), -1))
        p = int(round(e * float(np.exp(rng.normal(.05, .06))), -1))
        ws.append([sec, title, e, p, p - e, round(100 * (p - e) / e, 1)])
    wb.save(path)


def write_oews(path, x):
    rng = np.random.default_rng(W.SEED + 8)
    occ = [("11-1021", "General and Operations Managers", 48.1), ("13-2011", "Accountants and Auditors", 38.6),
           ("15-1232", "Computer User Support Specialists", 26.9), ("29-1141", "Registered Nurses", 41.2),
           ("29-2061", "Licensed Practical and Licensed Vocational Nurses", 27.4), ("31-1131", "Nursing Assistants", 18.3),
           ("31-9092", "Medical Assistants", 19.1), ("35-2014", "Cooks, Restaurant", 16.2),
           ("35-3023", "Fast Food and Counter Workers", 13.4), ("41-2031", "Retail Salespersons", 15.1),
           ("43-4051", "Customer Service Representatives", 18.9), ("47-2061", "Construction Laborers", 21.7),
           ("47-2111", "Electricians", 29.6), ("49-3023", "Automotive Service Technicians and Mechanics", 23.8),
           ("49-9071", "Maintenance and Repair Workers, General", 22.4), ("51-4041", "Machinists", 23.1),
           ("51-4121", "Welders, Cutters, Solderers, and Brazers", 22.6), ("53-3032", "Heavy and Tractor-Trailer Truck Drivers", 25.3),
           ("53-7062", "Laborers and Freight, Stock, and Material Movers, Hand", 17.8)]
    wb = Workbook()
    ws = wb.active
    ws.title = "Tallis County"
    ws.append(["Occupational Employment and Wage Statistics, Tallis County balance of state area, May 2025"])
    ws.append(["Hourly wages in dollars. Estimates are from a survey of establishments and are not comparable "
               "across years for trend analysis."])
    ws.append([])
    ws.append(["SOC code", "Occupation", "Employment", "Mean hourly wage", "Median hourly wage", "Annual mean wage"])
    for code, t, w in occ:
        mean = round(w * float(np.exp(rng.normal(0, .03))), 2)
        ws.append([code, t, int(round(float(rng.uniform(80, 1900)), -1)), mean,
                   round(mean * float(rng.uniform(.9, .98)), 2), int(round(mean * 2080, -1))])
    wb.save(path)
