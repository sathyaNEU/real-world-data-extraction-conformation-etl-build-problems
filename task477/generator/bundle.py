"""Reads the cut pack back: the county file's three releases conformed (the latest release of each quarter, both
layouts), Tallis County's private sectors as the routes in methods.py take them, and the two asks."""
import csv
import datetime as dt
import os

import world as W


def find(folder, stem):
    for fn in sorted(os.listdir(folder)):
        if fn.startswith(stem):
            return os.path.join(folder, fn)
    raise FileNotFoundError(stem)


def _rows(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def qcew(target, fips=W.COUNTY_FIPS, agg="74", own="5", use_latest=True):
    """{industry: {(y, q): (wages, (m1, m2, m3), disclosure)}} for one county, one aggregation level."""
    out = {}
    files = [W.FILES["qcew_a"], W.FILES["qcew_b"], W.FILES["qcew_c"]]
    if not use_latest:
        files = files[:2]
    for fn in files:
        for r in _rows(os.path.join(target, fn)):
            if r["area_fips"] != fips or r["agglvl_code"] != agg or r["own_code"] != own:
                continue
            k = (int(r["year"]), int(r["qtr"]))
            out.setdefault(r["industry_code"], {})[k] = (
                float(r["total_qtrly_wages"]), (int(r["month1_emplvl"]), int(r["month2_emplvl"]),
                                                int(r["month3_emplvl"])), r["disclosure_code"])
    return out


def record(target, **kw):
    """The routes' input: Tallis County's disclosed private sectors with every quarter 2013 to 2025."""
    q = qcew(target, **kw)
    rec = {}
    for ind, d in q.items():
        if len(d) == 4 * (W.Y1 - W.Y0 + 1) and all(v[2] == "" for v in d.values()):
            rec[ind] = {k: (v[0], v[1]) for k, v in d.items()}
    return rec


def _qafter(d, n):
    y, q = d.year, (d.month - 1) // 3 + 1 + n
    while q > 4:
        q -= 4
        y += 1
    return y, q


def ask_a(target, peo=True, last_exit=True, most_pay=True, merge=True):
    """Each industry's 2025 exiters working in it in the second quarter after exit."""
    emp = {r["ui_acct"]: r for r in _rows(find(target, "ui_employer_accounts"))}
    client = {}
    for r in _rows(find(target, "peo_client_registrations")):
        client[(r["peo_acct"], r["client_seq"])] = r["client_ui_acct"]
    last, who = {}, {}
    for r in _rows(find(target, "wioa_exits")):
        d = dt.datetime.strptime(r["exit_date"], "%m/%d/%Y").date()
        p = r["participant_id"]
        last[p] = max(last.get(p, d), d) if last_exit else min(last.get(p, d), d)
        who[p] = (r["birth_date"], r["home_zip"])
    # one person under two participant IDs: the same birth date and ZIP code and the same wage records
    sig = {}
    for r in _rows(find(target, "ui_wage_match_return")):
        sig.setdefault(r["participant_id"], []).append((r["ui_acct"], r["rptg_unit"], r["wage_yr"], r["wage_qtr"],
                                                         r["wages"]))
    alias = {}
    if merge:
        seen = {}
        for p in sorted(last):
            k = (who[p], tuple(sorted(sig.get(p, []))))
            if not k[1]:
                continue
            if k in seen:
                alias[p] = seen[k]
                last[seen[k]] = max(last[seen[k]], last[p])
            else:
                seen[k] = p
    pay = {}
    for r in _rows(find(target, "ui_wage_match_return")):
        p = r["participant_id"]
        if p in alias:
            continue
        if p not in last or (int(r["wage_yr"]), int(r["wage_qtr"])) != _qafter(last[p], 2):
            continue
        w = float(r["wages"])
        if w <= 0:
            continue
        acct = r["ui_acct"]
        if peo and (acct, r["rptg_unit"]) in client:
            acct = client[(acct, r["rptg_unit"])]
        pay.setdefault(p, {})
        pay[p][acct] = pay[p].get(acct, 0) + w
    out = {}
    for p, d in pay.items():
        accts = [max(d, key=lambda a: (d[a], a))] if most_pay else list(d)
        for a in accts:
            n6 = emp[a]["naics"]
            sec = {"31": "31-33", "32": "31-33", "33": "31-33", "44": "44-45", "45": "44-45", "48": "48-49",
                   "49": "48-49"}.get(n6[:2], n6[:2])
            out[sec] = out.get(sec, 0) + 1
    return out


def ask_b(target, route="checks"):
    """Each industry's 2025 individual training account payments: the checks issued in 2025 and not voided, each
    line's installment assigned to the industry of the program the line paid for (its title in the ledger, matched
    to the provider list's record of that program, active or withdrawn)."""
    from openpyxl import load_workbook
    wb = load_workbook(find(target, "etpl_programs"), read_only=True)
    ws = wb[wb.sheetnames[0]]
    hdr, sec, title = None, {}, {}
    for row in ws.iter_rows(values_only=True):
        if hdr is None:
            if row and row[0] == "Program ID":
                hdr = list(row)
            continue
        if row[0]:
            s_ = str(row[hdr.index("Target Industry (NAICS)")]).split()[0]
            sec[row[0]] = s_
            title[row[hdr.index("Program")]] = s_
    led = _rows(find(target, "ita_voucher_ledger"))
    reg = {r["check_no"]: r for r in _rows(find(target, "fiscal_agent_check_register"))}
    paid = {c for c, r in reg.items() if r["status"] != "VOID" and r["issue_date"].startswith("2025")}
    lines = [l for l in led if l["check_no"] in paid]

    def ind(l):
        if route in ("checks",):
            return title[l["description"].split(" - ", 1)[1]]
        return sec[l["program_id"]]

    def part(l):
        amt, n, i = float(l["voucher_amount"]), int(l["of_installments"]), int(l["installment"])
        return round(amt / n, 2) if i < n else round(amt - (n - 1) * round(amt / n, 2), 2)
    out = {}
    if route in ("checks", "id"):
        for l in lines:
            out[ind(l)] = out.get(ind(l), 0) + part(l)
    elif route == "lines":
        for l in lines:
            out[ind(l)] = out.get(ind(l), 0) + float(l["voucher_amount"])
    elif route == "share":
        by = {}
        for l in lines:
            by.setdefault(l["check_no"], []).append(l)
        for c, ls in by.items():
            tot = sum(float(l["voucher_amount"]) for l in ls)
            for l in ls:
                out[ind(l)] = out.get(ind(l), 0) + float(reg[c]["amount"]) * float(l["voucher_amount"]) / tot
    return out
