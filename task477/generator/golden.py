"""The golden deliverables: the application table (CSV, one row per eligible industry) and the slate memo for the board
(docx with a bar chart of every eligible industry's growth in pay per job), and generator/figures.json."""
import csv
import math
import io
import json
import os

import bundle as B
import docs as DC
import methods as M
import world as W

HERE = os.path.dirname(os.path.abspath(__file__))
SHORT = {"23": "Construction", "31-33": "Manufacturing", "42": "Wholesale trade", "44-45": "Retail trade",
         "48-49": "Transportation and warehousing", "51": "Information", "52": "Finance and insurance",
         "53": "Real estate and rental and leasing", "54": "Professional, scientific, and technical services",
         "56": "Administrative and support and waste services", "61": "Educational services",
         "62": "Health care and social assistance", "71": "Arts, entertainment, and recreation",
         "72": "Accommodation and food services", "81": "Other services"}


def figures(target):
    rec = B.record(target)
    g = M.route(rec, "golden")
    paid = M.route(rec, "paid")
    slate, nxt = M.slate(g)
    pslate, _ = M.slate(paid)
    a, b = B.ask_a(target), B.ask_b(target)
    el = M.eligible(rec)
    order = sorted(el, key=lambda s: -g[s])
    exp = {s: M.exposures(rec[s]) for s in el}
    jobs = {s: int(round(M.annual_jobs(rec[s], W.LATEST))) for s in el}
    return dict(
        call=dict(slate=slate, slate_names=[SHORT[s] for s in slate], next_in_line=nxt, next_name=SHORT[nxt]),
        eligible=el, n_eligible=len(el), order=order,
        growth={s: int(round(100 * g[s])) for s in el}, growth_unrounded={s: round(100 * g[s], 3) for s in el},
        paid_growth={s: int(round(100 * paid[s])) for s in el}, paid_slate=pslate,
        thursday_a={s: round(exp[s]["TA"], 3) for s in el},
        jobs_2025=jobs, names={s: SHORT[s] for s in el},
        asks=dict(exiters={s: int(a.get(s, 0)) for s in el}, ita={s: int(math.floor(b.get(s, 0.0) + .5)) for s in el},
                  exiters_total=sum(int(a.get(s, 0)) for s in el), ita_total=int(round(sum(b.get(s, 0.0) for s in el)))),
        min_jobs=W.MIN_JOBS, base=W.BASE, latest=W.LATEST,
        constants=dict(qd="%.2f" % M.QD, step=14, days=365, county=W.COUNTY_FIPS, own="5", agglvl="74"),
        sections=dict(DC.SECTIONS, industry="4.3"))


def write_table(path, f):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(["naics_sector", "industry", "pay_per_job_growth_2021_2025_pct", "on_slate",
                "exiters_employed_q2_after_exit", "ita_payments_2025_usd"])
    for s in f["order"]:
        w.writerow([s, SHORT[s], f["growth"][s], "Yes" if s in f["call"]["slate"] else "No",
                    f["asks"]["exiters"][s], f["asks"]["ita"][s]])
    open(path, "w", encoding="utf-8", newline="").write(buf.getvalue())


def chart_png(path, f):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 8, "svg.hashsalt": "t477"})
    order = f["order"]
    fig, ax = plt.subplots(figsize=(6.6, 3.9), dpi=150)
    ys = list(range(len(order)))[::-1]
    cols = ["#2f6f9f" if s in f["call"]["slate"] else "#b9c6d1" for s in order]
    vals = [f["growth_unrounded"][s] for s in order]
    ax.barh(ys, vals, color=cols, height=.66)
    for y, s, v in zip(ys, order, vals):
        ax.text(v + .2, y, "%d%%" % f["growth"][s], va="center", fontsize=7.5)
    ax.set_yticks(ys)
    ax.set_yticklabels([SHORT[s] for s in order], fontsize=7.5)
    ax.set_xlim(0, max(vals) + 3)
    ax.set_xlabel("Growth in pay per job, 2021 to 2025 (pay for the work done in each year), per cent")
    ax.spines[["top", "right"]].set_visible(False)
    ax.set_title("Tallis County eligible industries: growth in pay per job, 2021 to 2025\nDark bars: proposed slate",
                 fontsize=8.5, loc="left")
    fig.tight_layout()
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)


def write_memo(path, f, png):
    from docx import Document
    from docx.shared import Inches, Pt
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    doc.add_heading("2027-2029 Sector Partnership Slate", level=1)
    for k, v in (("To", "Board of Directors, %s" % W.BOARD), ("From", "%s, %s" % W.PEOPLE["analyst"]),
                 ("Date", W.MEMO_DATE), ("Re", "Industry slate for the %s vote" % W.VOTE_DATE)):
        p = doc.add_paragraph()
        p.add_run("%s: " % k).bold = True
        p.add_run(v)
    s = f["call"]["slate_names"]
    doc.add_paragraph(
        "I recommend that the board name %s, %s, %s and %s as our slate for the %s 2027-2029 round. They are the "
        "four of our %d eligible industries whose pay per job grew fastest from 2021 to 2025. %s is next in line."
        % (s[0], s[1], s[2], s[3], W.FUND, f["n_eligible"], f["call"]["next_name"]))
    p62, p48 = f["paid_growth"]["62"], f["paid_growth"]["48-49"]
    same = f["growth"]["62"] == f["growth"]["48-49"]
    doc.add_paragraph(
        "The Fund measures pay for the work done in each year, but the state's county file counts wages in the quarter "
        "they are paid. New Year's Day 2026 fell on a Thursday and was paid on 31 December 2025, which gave weekly "
        "Thursday payrolls 53 pay dates in 2025 and half of the every-other-Thursday payrolls 27. Nothing like that "
        "happened in 2021. I estimated how much of each industry's payroll runs on each weekly and every-other-week "
        "schedule from its quarterly wages since 2013 and took the extra pay date out of 2025.")
    doc.add_paragraph(
        "It matters for two industries. Our hospitals and the large trucking and warehousing employers pay mostly on "
        "Thursdays, so on wages as paid %s shows %d%% and %s %d%%, enough to push %s and %s off the slate. For the work "
        "done %s." % (SHORT["62"], p62, SHORT["48-49"], p48, s[2], s[3],
                      "both grew %d%%" % f["growth"]["62"] if same else "they grew %d%% and %d%%" % (
                          f["growth"]["62"], f["growth"]["48-49"])))
    doc.add_picture(png, width=Inches(6.3))
    doc.add_paragraph("Application figures by industry (guidelines 3.2 and 3.3):")
    t = doc.add_table(rows=1, cols=5)
    t.style = "Light Grid Accent 1"
    for c, h in zip(t.rows[0].cells, ("Industry (NAICS)", "Pay per job growth", "Slate",
                                      "2025 exiters employed, 2nd quarter after exit", "2025 ITA payments")):
        c.text = h
    for sec in f["order"]:
        r = t.add_row().cells
        r[0].text = "%s (%s)" % (SHORT[sec], sec)
        r[1].text = "%d%%" % f["growth"][sec]
        r[2].text = "Yes" if sec in f["call"]["slate"] else ""
        r[3].text = "{:,}".format(f["asks"]["exiters"][sec])
        r[4].text = "${:,}".format(f["asks"]["ita"][sec])
    doc.add_paragraph(
        "Exiters are counted once, at their last 2025 exit (a customer registered twice is one exiter), in the industry of the business that paid them most in the "
        "second quarter after exit; workers leased through a professional employer organization are counted in the "
        "client business's industry. ITA payments are the fiscal agent's 2025 checks net of voids, each check's "
        "installments assigned to the industry of the program they paid for.")
    doc.core_properties.author = W.PEOPLE["analyst"][0]
    doc.save(path)


def build(gdir, x, target=None, figs=None):
    import postprocess as PP
    target = target or os.path.join(os.path.dirname(gdir), "target")
    figs = figs or os.path.join(HERE, "figures.json")
    os.makedirs(gdir, exist_ok=True)
    f = figures(target)
    json.dump(f, open(figs, "w", encoding="utf-8"), indent=1, sort_keys=True)
    write_table(os.path.join(gdir, W.GOLDEN["table"]), f)
    png = os.path.join(gdir, ".chart.png")
    chart_png(png, f)
    write_memo(os.path.join(gdir, W.GOLDEN["memo"]), f, png)
    os.remove(png)
    PP.normalize_ooxml(os.path.join(gdir, W.GOLDEN["memo"]), W.PEOPLE["analyst"][0], W.PEOPLE["analyst"][0],
                       "2026-10-13T20:11:00Z", "2026-10-14T14:37:00Z", company=W.BOARD)
    return f
