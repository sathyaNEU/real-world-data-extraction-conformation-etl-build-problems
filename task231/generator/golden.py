"""The two deliverables: bcs_sep2026_extension.xlsx (the call with every region's alert level, the contact centre's
work so far, the June 2026 late returns and the March 2026 refreshment) and bcs_sep2026_response.pptx (each region's
response rate against the 55 and 58 per cent lines)."""
import copy
import datetime as _dt
import io
import os
import re
import zipfile

from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

import postprocess
import world as W

WORKBOOK = "bcs_sep2026_extension.xlsx"
SLIDE = "bcs_sep2026_response.pptx"
LEVEL = {"red": "Red", "amber": "Amber", "green": "Green"}
ACTION = {"red": "Reminder-call extension, 28 Sep to 9 Oct", "amber": "Extra email reminder, 28 Sep", "green": "None"}
COLOUR = {"red": "B3261E", "amber": "D98E04", "green": "3C7A3E"}
THIN = Side(style="thin", color="B7BEC8")


def _header(ws, row, cols):
    for j, h in enumerate(cols, start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="31506F")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = Border(bottom=THIN)


def _title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = sub
    ws["A2"].font = Font(italic=True, size=9, color="555555")


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _top(ws, title, asof, ncols):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = asof
    ws["A2"].font = Font(italic=True, size=9, color="555555")
    ws.freeze_panes = "A5"
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _total(ws, row, values, fmts):
    for j, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=j, value=v)
        c.font = Font(bold=True)
        c.border = Border(top=THIN)
        if j in fmts:
            c.number_format = fmts[j]


def _note(ws, row, text, widths):
    """a note wrapped across the table's columns, its row tall enough for the text"""
    last = get_column_letter(len(widths))
    ws.merge_cells("A%d:%s%d" % (row, last, row))
    c = ws.cell(row=row, column=1, value=text)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    c.font = Font(size=9, color="333333")
    chars = sum(widths) * 1.15
    ws.row_dimensions[row].height = 13.5 * max(1, -(-len(text) // int(chars)))


ASOF = "BCS September 2026 round, as at the main close, 17:00 on 25 September 2026. Business Surveys Unit, 28 September 2026."


def workbook(A, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Alert levels"
    _top(ws, "Alert levels at the main close: Capital and Harbour into the extension, Islands and Western Uplands an extra "
         "email", ASOF, 6)
    _header(ws, 4, ["Region", "Region name", "Returned by the main close", "Response rate (%)", "Alert level",
                    "Further fieldwork"])
    r = 5
    for g in W.REGS:
        lv = A["levels"][g]
        vals = [g, W.NAME[g], A["R"][g], round(A["rate"][g], 1), LEVEL[lv], ACTION[lv]]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.border = Border(bottom=THIN)
            if j == 3:
                c.number_format = "#,##0"
            if j == 4:
                c.number_format = "0.0"
            if j == 5:
                c.font = Font(bold=True, color=COLOUR[lv])
        r += 1
    _total(ws, r, ["All regions", "", sum(A["R"].values())], {3: "#,##0"})
    r += 2
    notes = [
        "Extension (red): %s. Extra email (amber): %s. Neither (green): the other %d regions."
        % (" and ".join(W.NAME[g] for g in A["extension"]), " and ".join(W.NAME[g] for g in A["email"]), len(A["neither"])),
        "Response rate: returns by 17:00 on 25 September over the invited firms still trading on 1 September (standard v3.0, "
        "section 4). Firms the business register already shows as having ceased by 1 September are out of scope; the other "
        "non-returners are counted in scope at %.1f per cent, the share of earlier rounds' non-returners (March 2017 to "
        "September 2022) that the register has not since shown as ceased by their round's frame date." % A["inscope_pct"],
        "Levels per standard v3.0, section 5: red under 55, amber 55 to under 58, green 58 or more. Rates rounded for "
        "display; levels set on the unrounded rates.",
    ]
    widths = [12, 18, 14, 13, 11, 38]
    for line in notes:
        _note(ws, r, line, widths)
        r += 1
    _widths(ws, widths)

    ws = wb.create_sheet("Contact centre")
    _top(ws, "Contact centre, September round so far, by region", ASOF, 4)
    _header(ws, 4, ["Region", "Region name", "Firms reached so far", "Calls placed so far"])
    for i, g in enumerate(W.REGS, start=5):
        for j, v in enumerate([g, W.NAME[g], A["reached"][g], A["calls"][g]], start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = Border(bottom=THIN)
            if j >= 3:
                c.number_format = "#,##0"
    tot = len(W.REGS) + 5
    _total(ws, tot, ["All regions", "", sum(A["reached"].values()), sum(A["calls"].values())], {3: "#,##0", 4: "#,##0"})
    widths = [12, 18, 14, 14]
    _note(ws, tot + 2, "Both teams, to the 17:00 close on 25 September. Calls placed counts each call the centre placed once: "
          "the contractor's records for 22 September were loaded a second time on 23 September, and a call passed to a "
          "supervisor (code 71) continues in the supervisor's record. Calls firms made to the centre are not calls placed.",
          widths)
    _note(ws, tot + 3, "Reached: the centre has spoken with the respondent or someone else at the firm, on either kind of call "
          "(contact centre outcome codes).", widths)
    _widths(ws, widths)

    ws = wb.create_sheet("June late returns")
    _top(ws, "June 2026 round: returns that reached us after the main close, by region", ASOF, 5)
    _header(ws, 4, ["Region", "Region name", "June returns", "Reached us after the main close", "Share of June returns (%)"])
    for i, g in enumerate(W.REGS, start=5):
        for j, v in enumerate([g, W.NAME[g], A["june_returns"][g], A["late"][g], A["late_share"][g]], start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = Border(bottom=THIN)
            if j == 5:
                c.number_format = "0.0"
    tot = len(W.REGS) + 5
    jr, lt = sum(A["june_returns"].values()), sum(A["late"].values())
    _total(ws, tot, ["All regions", "", jr, lt, round(100.0 * lt / jr, 1)], {3: "#,##0", 5: "0.0"})
    widths = [12, 18, 12, 16, 16]
    _note(ws, tot + 2, "Main close 17:00 on 19 June. A return counts from the moment it reached the bank: web submission, post "
          "room log, or the end of the telephone interview in the contact centre's records (the survey database loads "
          "telephone returns overnight). Each firm's return counted once.", widths)
    _widths(ws, widths)

    ws = wb.create_sheet("March 2026 recruits")
    _top(ws, "March 2026 refreshment: firms recruited and firms that have returned at least once, by industry group", ASOF, 4)
    _header(ws, 4, ["Industry group", "Group name", "Recruited", "Returned at least once"])
    for i, (g, name) in enumerate(W.GROUPS, start=5):
        for j, v in enumerate([g, name, A["recruited"][g], A["returned"][g]], start=1):
            ws.cell(row=i, column=j, value=v).border = Border(bottom=THIN)
    tot = len(W.GROUPS) + 5
    _total(ws, tot, ["All groups", "", sum(A["recruited"].values()), sum(A["returned"].values())], {})
    widths = [14, 38, 11, 20]
    _note(ws, tot + 2, "Firms recruited in the March 2026 refreshment, counted once each, including those that withdrew before "
          "their first round; NACE Rev. 2.1 sections mapped to the survey's groups. Returns are each firm's own, in the March "
          "and June 2026 rounds and the September round to the close; returns logged under a reissued panel ID before the "
          "firm joined belong to the earlier holder.", widths)
    _widths(ws, widths)

    ws = wb.create_sheet("Basis")
    _top(ws, "Basis of the response rates", ASOF, 2)
    rows = [
        ["Settled rounds", "March 2017 to September 2022 (frame dates at least four years before the extract)"],
        ["Non-returners at the main close, not then shown closed", A["settled_den"]],
        ["Of them, ceased by the round's frame date per the register now", A["settled_num"]],
        ["Closure share", "%.1f%%" % A["settled_pct"]],
        ["September invited", A["invited_total"]],
        ["September returned by the main close", A["R_total"]],
        ["Ceased by 1 September per the register", A["known_total"]],
        ["Other non-returners", A["n_total"]],
    ]
    _header(ws, 4, ["Item", "Value"])
    for i, rr in enumerate(rows, start=5):
        ws.cell(row=i, column=1, value=rr[0])
        c = ws.cell(row=i, column=2, value=rr[1])
        if isinstance(rr[1], int):
            c.number_format = "#,##0"
        c.alignment = Alignment(horizontal="right")
    widths = [58, 70]
    _note(ws, len(rows) + 6, "The register records a closure when it is notified, often a year or more after trading stops, so "
          "it does not yet show many of the closures among this round's non-returners; the share comes from rounds old enough "
          "for the register to have caught up.", widths)
    _widths(ws, widths)
    wb.properties.creator = W.PEOPLE["analyst"]
    wb.properties.lastModifiedBy = W.PEOPLE["analyst"]
    wb.save(path)
    postprocess.scrub(path, W.PEOPLE["analyst"], "2026-09-28")
    postprocess.finish_office(path, "2026-09-26T08:12:00Z", "2026-09-28T06:41:00Z", W.ORG)


# ------------------------------------------------------------------------------------------------ slide
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _q(tag):
    p, t = tag.split(":")
    return "{%s}%s" % ({"c": C_NS, "a": A_NS}[p], t)


def _lines_to_line_chart(chart):
    """move the two threshold series out of the column chart into a line chart on the same axes"""
    cs = chart._chartSpace
    plot = cs.find(".//" + _q("c:plotArea"))
    bar = plot.find(_q("c:barChart"))
    sers = bar.findall(_q("c:ser"))
    line = etree.SubElement(plot, _q("c:lineChart"))
    plot.remove(line)
    bar.addnext(line)
    etree.SubElement(line, _q("c:grouping")).set("val", "standard")
    etree.SubElement(line, _q("c:varyColors")).set("val", "0")
    for i, s in enumerate(sers[1:]):
        bar.remove(s)
        for tag in ("c:invertIfNegative",):
            e = s.find(_q(tag))
            if e is not None:
                s.remove(e)
        sp = s.find(_q("c:spPr"))
        if sp is None:
            sp = etree.Element(_q("c:spPr"))
            s.insert(3, sp)
        for ch in list(sp):
            sp.remove(ch)
        ln = etree.SubElement(sp, _q("a:ln"))
        ln.set("w", "22225")
        sf = etree.SubElement(ln, _q("a:solidFill"))
        etree.SubElement(sf, _q("a:srgbClr")).set("val", ["B3261E", "D98E04"][i])
        etree.SubElement(ln, _q("a:prstDash")).set("val", "dash")
        mk = etree.Element(_q("c:marker"))
        etree.SubElement(mk, _q("c:symbol")).set("val", "none")
        sp.addnext(mk)
        etree.SubElement(s, _q("c:smooth")).set("val", "0")
        line.append(s)
    etree.SubElement(line, _q("c:marker")).set("val", "1")
    for ax in bar.findall(_q("c:axId")):
        etree.SubElement(line, _q("c:axId")).set("val", ax.get("val"))


def slide(A, path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[5])
    ext = [W.NAME[g] for g in A["extension"]]
    em = [W.NAME[g] for g in A["email"]]
    s.shapes.title.text = "%s go into the reminder-call extension; %s get the extra email" % (" and ".join(ext), " and ".join(em))
    ttl = s.shapes.title
    ttl.left, ttl.top, ttl.width, ttl.height = Inches(0.4), Inches(0.3), Inches(12.5), Inches(0.95)
    tf = ttl.text_frame
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    tf.paragraphs[0].runs[0].font.size = Pt(26)
    tf.paragraphs[0].runs[0].font.bold = True
    order = sorted(W.REGS, key=lambda g: (A["rate"][g], g))
    cd = CategoryChartData()
    cd.categories = [W.NAME[g] for g in order]
    cd.add_series("Response rate at the main close (%)", [round(A["rate"][g], 1) for g in order])
    cd.add_series("55%: extension below", [55.0] * len(order))
    cd.add_series("58%: extra email below", [58.0] * len(order))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.4), Inches(1.3), Inches(12.5), Inches(5.45), cd)
    ch = gf.chart
    ch.has_legend = True
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(10)
    from pptx.enum.chart import XL_LEGEND_POSITION
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.has_title = True
    ch.chart_title.text_frame.text = ("Response rate at the 17:00 main close on 25 September 2026, by region (%): red under 55, "
                                      "amber 55 to under 58, green 58 or more")
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    va = ch.value_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = 40, 65, 5
    va.has_major_gridlines = False
    va.tick_labels.font.size = Pt(10)
    ch.category_axis.tick_labels.font.size = Pt(10)
    pl = ch.plots[0]
    pl.gap_width = 60
    ser = pl.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = RGBColor(0x3C, 0x7A, 0x3E)
    for i, g in enumerate(order):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = RGBColor.from_string(COLOUR[A["levels"][g]])
    ser.data_labels.show_value = True
    ser.data_labels.number_format = "0.0"
    ser.data_labels.number_format_is_linked = False
    ser.data_labels.font.size = Pt(10)
    ser.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    _lines_to_line_chart(ch)
    leg = ch._chartSpace.find(".//" + _q("c:legend"))
    le = etree.Element(_q("c:legendEntry"))
    etree.SubElement(le, _q("c:idx")).set("val", "0")
    etree.SubElement(le, _q("c:delete")).set("val", "1")
    leg.insert(1, le)
    tb = s.shapes.add_textbox(Inches(0.4), Inches(6.78), Inches(12.3), Inches(0.4))
    tb.text_frame.word_wrap = True
    tb.text_frame.text = ("Returns by the close over invited firms still trading on 1 September. Non-returners without a "
                          "closure on the business register counted in scope at %.1f%%, from rounds March 2017 to September "
                          "2022 that the register has caught up with." % A["inscope_pct"])
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
    tb.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    src = s.shapes.add_textbox(Inches(0.4), Inches(7.12), Inches(11.8), Inches(0.3))
    src.text_frame.text = ("Source: BCS fieldwork system, survey database and Brevland Business Register extract of 25 September "
                           "2026. Business Surveys Unit, 28 September 2026.")
    src.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
    src.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    sn = s.shapes.add_textbox(Inches(12.55), Inches(7.12), Inches(0.4), Inches(0.3))
    sn.text_frame.text = "1"
    sn.text_frame.paragraphs[0].runs[0].font.size = Pt(8)
    sn.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x77, 0x77, 0x77)
    prs.core_properties.author = W.PEOPLE["analyst"]
    prs.core_properties.last_modified_by = W.PEOPLE["analyst"]
    prs.core_properties.title = "BCS September 2026: response at the main close"
    prs.core_properties.comments = ""
    prs.core_properties.created = _dt.datetime(2026, 9, 26, 9, 5)
    prs.core_properties.modified = _dt.datetime(2026, 9, 28, 6, 52)
    prs.core_properties.revision = 4
    prs.save(path)
    _fix_embedded(path)
    postprocess.scrub(path, W.PEOPLE["analyst"], "2026-09-28")


def _fix_embedded(path):
    """the chart workbook python-pptx embeds carries the build clock: give it the deck's own date and fixed entry stamps so
    two builds are byte-identical"""
    with zipfile.ZipFile(path) as z:
        items = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    out = []
    for n, b in items:
        if n.endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(b)) as zz:
                inner = [(i.filename, zz.read(i.filename)) for i in zz.infolist()]
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zo:
                for m, bb in inner:
                    if m == "docProps/core.xml":
                        bb = re.sub(rb"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", b"2026-09-28T06:52:00Z", bb)
                        bb = re.sub(rb"<dc:creator>[^<]*</dc:creator>", b"<dc:creator>" + W.PEOPLE["analyst"].encode()
                                    + b"</dc:creator>", bb)
                    zi = zipfile.ZipInfo(m, date_time=(1980, 1, 1, 0, 0, 0))
                    zi.compress_type = zipfile.ZIP_DEFLATED
                    zo.writestr(zi, bb)
            b = buf.getvalue()
        if n.startswith("ppt/slideMasters/") and n.endswith(".xml"):
            b = b.replace("\u2013".encode(), "\u2022".encode())
        out.append((n, b))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in out:
            zi = zipfile.ZipInfo(n, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, b)


def write(A, golden_dir):
    os.makedirs(golden_dir, exist_ok=True)
    workbook(A, os.path.join(golden_dir, WORKBOOK))
    slide(A, os.path.join(golden_dir, SLIDE))
