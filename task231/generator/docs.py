"""task231 documents: the fieldwork quality standard, the register's notes on its closure extract, the questionnaire,
the round calendar, the contact centre's outcome codes, the data dictionary and the published June 2026 balances."""
from datetime import date, timedelta

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import postprocess
import records
import sim
import world as W


# ------------------------------------------------------------------------------------------------ helpers
def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=13, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Helvetica-Bold", fontSize=14, leading=18, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=14, spaceBefore=8,
                        spaceAfter=4)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10, textColor=colors.HexColor("#444444"))
    item = ParagraphStyle("item", parent=body, leftIndent=14, spaceAfter=3)
    return dict(body=body, h1=h1, h2=h2, small=small, item=item)


def _footer(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(18 * mm, 10 * mm, text)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _pdf(path, story, title, author, footer, stamp):
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=18 * mm, title=title, author=author, creator=author, invariant=1)
    doc.build(story, onFirstPage=_footer(footer), onLaterPages=_footer(footer))
    postprocess.scrub(path, author, stamp)


def _table(rows, widths, size=8.5):
    t = Table(rows, colWidths=widths)
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", size), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", size),
                           ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E4E8EE")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    return t


HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="31506F")


def _sheet(ws, header, rows, widths=None, wrap=False):
    ws.append(header)
    for c in ws[1]:
        c.font = HEAD
        c.fill = FILL
        c.alignment = Alignment(vertical="top", wrap_text=True)
    for r in rows:
        ws.append(r)
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    if wrap:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


def _save(wb, path, who, created, modified):
    wb.properties.creator = who
    wb.properties.lastModifiedBy = who
    wb.save(path)
    postprocess.scrub(path, who, modified[:10])
    postprocess.finish_office(path, created, modified, W.ORG)


# ------------------------------------------------------------------------------------------------ the standard
def standard(path):
    S = _styles()
    st = []
    st.append(Paragraph("Business Conditions Survey: fieldwork quality standard", S["h1"]))
    st.append(Paragraph("%s, %s. Version 3.0, approved 6 July 2026; applies from the September 2026 round. "
                        "Owner: Head of Business Surveys (%s). Approver: Director of Economic Statistics (%s)."
                        % (W.ORG, W.UNIT, W.PEOPLE["head"], W.PEOPLE["director"]), S["small"]))
    st.append(Spacer(1, 6))
    sec = [
        ("1 Purpose", [
            "This standard sets how the Business Surveys Unit follows fieldwork in each round of the Business Conditions "
            "Survey (BCS), and what further fieldwork a region gets when its response falls short at the main close. It "
            "replaces version 2.1 of March 2021, under which every region received the same reminders."]),
        ("2 The round", [
            "2.1 The BCS is a quarterly survey of a panel of firms with five or more employees in the bank's twelve regions. "
            "Rounds run in March, June, September and December.",
            "2.2 The frame date of a round is the first day of the round month. Every firm on the panel register on the frame "
            "date is invited to the round.",
            "2.3 Invited firms that have not returned the questionnaire receive an email reminder and, from the date in the "
            "round calendar, reminder calls from the bank's contact centre and, on its busiest days, from its overflow "
            "contractor.",
            "2.4 The main close is 17:00 Brevland time on the date the round calendar (BCS_round_calendar.xlsx) gives. Returns "
            "that reach the bank after the main close are used in the round's results if they arrive by the data cut-off."]),
        ("3 Records", [
            "3.1 The panel register lists every firm recruited to the panel, when it was recruited and, if it has left, when "
            "and why.",
            "3.2 The survey database holds the outcome of every invited firm in every round.",
            "3.3 The fieldwork system logs each return of the current round as it reaches the bank.",
            "3.4 The contact centre's system holds the record of every reminder call, the bank's own centre's and the "
            "contractor's.",
            "3.5 The Brevland Business Register sends, each quarter and on request, the closure records it holds for the "
            "company numbers on the panel register."]),
        ("4 Response rate", [
            "4.1 A region's response rate for a round is the number of panel firms in the region that returned the "
            "questionnaire by the main close, divided by the number of panel firms in the region that were invited to the "
            "round and were trading on the round's frame date, as a percentage.",
            "4.2 A firm whose business ceased trading on or before the frame date is out of scope for the round, whether or "
            "not the Brevland Business Register yet records its closure. A firm that was trading on the frame date stays in "
            "scope for the round whatever happens to it afterwards.",
            "4.3 A firm has returned the questionnaire from the moment its first return reaches the bank: when it submits "
            "the web questionnaire, when its telephone interview with the contact centre ends, or when the post room logs "
            "its paper questionnaire. A firm that returns more than once counts once."]),
        ("5 Alert levels", [
            "5.1 At the main close the unit sets each region's alert level from its response rate for the round:"]),
    ]
    for head, paras in sec:
        st.append(Paragraph(head, S["h2"]))
        for p in paras:
            st.append(Paragraph(p, S["body"]))
    st.append(_table([["Alert level", "Response rate", "Further fieldwork in the round"],
                      ["Red", "under 55 per cent", "The region goes into the reminder-call extension (section 6)."],
                      ["Amber", "55 per cent or more and under 58 per cent",
                       Paragraph("Firms in the region that have not returned receive one extra email reminder on the first "
                                 "working day after the main close.", S["body"])],
                      ["Green", "58 per cent or more", "None."]], [24 * mm, 60 * mm, 90 * mm]))
    st.append(Spacer(1, 4))
    st.append(Paragraph("5.2 Alert levels are set once for each round and are not revised during the extension.", S["body"]))
    for head, paras in [
        ("6 The reminder-call extension", [
            "6.1 In red regions the contact centre calls the firms that have not returned for two further weeks from the "
            "Monday after the main close; for the September 2026 round, 28 September to 9 October 2026.",
            "6.2 The contact centre plans the extension's call lists and its contractor's share of them from its own "
            "records of the round so far.",
            "6.3 Returns that arrive during the extension count towards the round's results; they do not change the alert "
            "level."]),
        ("7 Reporting", [
            "7.1 On the first working day after the main close the unit sends the Director of Economic Statistics each "
            "region's alert level for the round, with the contact centre's work in the round so far."]),
    ]:
        st.append(Paragraph(head, S["h2"]))
        for p in paras:
            st.append(Paragraph(p, S["body"]))
    _pdf(path, st, "BCS fieldwork quality standard, version 3.0", W.PEOPLE["head"],
         "BCS fieldwork quality standard v3.0  |  %s, Business Surveys Unit" % W.ORG, "2026-07-06")


# ------------------------------------------------------------------------------------------------ register notes
def register_notes(path):
    S = _styles()
    st = [Paragraph("Closure records for Brevland Central Bank: extract of 25 September 2026", S["h1"]),
          Paragraph("%s, data services. Supplied to the bank's Business Surveys Unit under the standing data agreement of "
                    "January 2014." % W.REGISTER, S["small"]), Spacer(1, 6)]
    for head, paras in [
        ("1 What the extract holds", [
            "The extract holds the closure records on the register at the end of 25 September 2026 for the company numbers "
            "on the bank's panel register, as the bank supplied them on 21 September 2026. A company appears in the extract "
            "only if its closure has been registered. A company that has stopped trading but whose closure has not been "
            "registered does not appear."]),
        ("2 How a closure is registered", [
            "A closure is registered when it is notified to the register or when the register's own procedure for the "
            "company is complete. There are three procedures."]),
    ]:
        st.append(Paragraph(head, S["h2"]))
        for p in paras:
            st.append(Paragraph(p, S["body"]))
    for p in ["<b>Voluntary dissolution.</b> The directors of a company that has stopped trading apply to dissolve it; the "
              "dissolution is registered when the register accepts the application.",
              "<b>Administrative strike-off.</b> The register strikes off a company that has stopped filing its annual "
              "accounts and returns, at the end of its strike-off procedure.",
              "<b>Liquidation.</b> A company wound up by a liquidator is registered as dissolved when the liquidator's final "
              "notice is filed."]:
        st.append(Paragraph(p, S["item"]))
    st.append(Paragraph("3 Fields", S["h2"]))
    st.append(_table([["Field", "Content"],
                      ["company_number", Paragraph("The company number in the register's current form: six digits, a hyphen "
                                                   "and a check digit. Numbers issued before 1 March 2021 keep their six "
                                                   "digits; the check digit was added to them on that date.", S["body"])],
                      ["ceased_trading", Paragraph("The date the company stopped trading, as established in its closure "
                                                   "procedure.", S["body"])],
                      ["registered", Paragraph("The date the closure was entered in the register.", S["body"])],
                      ["procedure", Paragraph("Voluntary dissolution, administrative strike-off or liquidation.", S["body"])]],
                     [32 * mm, 142 * mm]))
    st.append(Paragraph("4 Use", S["h2"]))
    st.append(Paragraph("The records are released to the bank for statistical work under the data agreement and may not be "
                        "passed on. Questions about the extract go to the register's data services desk.", S["body"]))
    _pdf(path, st, "Closure records for Brevland Central Bank, extract of 25 September 2026", W.REGISTER,
         "%s  |  data services  |  closure extract notes" % W.REGISTER, "2026-09-25")


# ------------------------------------------------------------------------------------------------ questionnaire (distractor)
def questionnaire(path):
    S = _styles()
    st = [Paragraph("Business Conditions Survey 2026", S["h1"]),
          Paragraph("%s. Please answer for your business in Brevland as a whole. Return the questionnaire on the web "
                    "(the link in your invitation email), by post in the reply envelope, or by telephone with our contact "
                    "centre." % W.ORG, S["small"]), Spacer(1, 6)]
    qs = [("1", "Compared with the previous three months, has your output over the past three months gone up, stayed the "
                "same or gone down?"),
          ("2", "Over the next three months, do you expect your output to go up, stay the same or go down?"),
          ("3", "Compared with the previous three months, have new orders over the past three months gone up, stayed the "
                "same or gone down?"),
          ("4", "Over the next three months, do you expect the number of people you employ to go up, stay the same or go "
                "down?"),
          ("5", "Over the next three months, do you expect your average selling prices to go up, stay the same or go down?"),
          ("6", "At what percentage of full capacity is your business working at present?"),
          ("7", "Which of these, if any, is limiting your activity at present: demand, shortage of labour, shortage of "
                "materials or equipment, access to finance, other (please say)?"),
          ("8", "How would you describe your business's overall situation at present: good, satisfactory or poor?")]
    st.append(Paragraph("About your business", S["h2"]))
    st.append(_table([["Item", "Please give"],
                      ["A", Paragraph("Number of people employed at the end of last month, including working owners.", S["body"])],
                      ["B", Paragraph("Main activity, in a few words, if it has changed since your last return.", S["body"])],
                      ["C", Paragraph("Name, telephone number and email address of the person completing this questionnaire.",
                                      S["body"])]], [20 * mm, 154 * mm]))
    st.append(Paragraph("Business conditions", S["h2"]))
    st.append(_table([["Question", "Text"]] + [[q, Paragraph(t, S["body"])] for q, t in qs], [20 * mm, 154 * mm]))
    st.append(Paragraph("Returning the questionnaire", S["h2"]))
    for t in ["On the web: use the link in your invitation email; your answers are saved as you go and sent when you submit.",
              "By post: use the reply envelope; no stamp is needed.",
              "By telephone: our contact centre can take your answers by telephone. If we have left you a message, call the "
              "number in the message on a working day between 08:30 and 17:30.",
              "Please return the questionnaire by the date in your invitation. Your answers are used only for the bank's "
              "statistics and are published only as totals in which no business can be identified."]:
        st.append(Paragraph(t, S["body"]))
    _pdf(path, st, "Business Conditions Survey questionnaire 2026", W.ORG, "BCS questionnaire 2026  |  %s" % W.ORG,
         "2026-02-10")


# ------------------------------------------------------------------------------------------------ calendar
def _wd(d):
    while d.weekday() >= 5:
        d += timedelta(days=1)
    return d


def calendar_rows():
    rows = []
    frames = [f for f in records.FRAMES if f >= P.FIRST_ROUND] + [date(2026, 12, 1)]
    for fr in frames:
        close = sim.close_of(fr)
        days = sim.call_days(fr, close)
        cut = sim.end_of(fr)
        pub = cut
        n = 0
        while n < 10:
            pub += timedelta(days=1)
            if pub.weekday() < 5:
                n += 1
        email = _wd(fr + timedelta(days=7))
        extra = ext = ""
        if fr >= P.CURRENT:
            extra = _wd(close + timedelta(days=1)).isoformat()
            mon = close + timedelta(days=3)
            ext = "%s to %s" % (mon.isoformat(), (mon + timedelta(days=11)).isoformat())
        rows.append([W.round_code(fr), W.round_quarter(fr), fr.isoformat(), fr.isoformat(), email.isoformat(),
                     (fr + timedelta(days=days[0])).isoformat(), close.isoformat(), "17:00", cut.isoformat(), pub.isoformat(),
                     extra, ext, "planned" if fr > P.CURRENT else ""])
    return rows


CAL_HEADER = ["round", "quarter", "frame_date", "questionnaire_sent", "email_reminder", "reminder_calls_from", "main_close",
              "main_close_time", "data_cut_off", "results_published", "extra_email_amber", "extension_calls_red", "status"]


def calendar(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Rounds"
    _sheet(ws, CAL_HEADER, calendar_rows(), widths=[9, 9, 12, 13, 13, 14, 12, 10, 12, 13, 14, 24, 9])
    ws2 = wb.create_sheet("Notes")
    for line in ["Round calendar of the Business Conditions Survey, March 2017 to December 2026.",
                 "Times are Brevland time. The main close is at 17:00 on the date shown.",
                 "The extra email and the extension apply from the September 2026 round (fieldwork quality standard v3.0, "
                 "section 5); earlier rounds ran under version 2.1, with the same reminders for every region.",
                 "Maintained by the Business Surveys Unit; last updated 14 September 2026."]:
        ws2.append([line])
    ws2.column_dimensions["A"].width = 120
    _save(wb, path, W.PEOPLE["analyst"], "2017-01-19T09:12:00Z", "2026-09-14T10:41:00Z")


# ------------------------------------------------------------------------------------------------ outcome codes
BCB_CODES = [("COMP", "Interview completed by telephone", "Y"),
             ("PROM", "Respondent will return the questionnaire", "Y"),
             ("REFU", "Respondent declined to take part this round", "Y"),
             ("STAF", "Answered by a colleague or reception; respondent not available; message left", "Y"),
             ("NOAN", "No answer", "N"),
             ("VMSG", "Voicemail; message left", "N"),
             ("BUSY", "Line busy", "N"),
             ("UNOB", "Number not in service", "N"),
             ("CEAS", "Told that the business has stopped trading", "N")]
OVF_CODES = [("10", "Interview completed", "COMP"), ("21", "Will return questionnaire", "PROM"),
             ("23", "Declined", "REFU"), ("31", "Answered, contact unavailable, message taken", "STAF"),
             ("41", "No answer", "NOAN"), ("42", "Answer machine, message left", "VMSG"), ("43", "Busy", "BUSY"),
             ("51", "Number unobtainable", "UNOB"), ("61", "Business closed", "CEAS"),
             ("71", "Transferred to supervisor", "")]


def codebook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "BCB codes"
    _sheet(ws, ["code", "outcome", "reached"], [list(x) for x in BCB_CODES], widths=[8, 80, 9])
    ws2 = wb.create_sheet("OVF codes")
    _sheet(ws2, ["code", "outcome (contractor's wording)", "BCB code"], [list(x) for x in OVF_CODES], widths=[8, 50, 10])
    ws3 = wb.create_sheet("Notes")
    for line in [
        "Outcome codes in the contact centre's call records (contact_centre_calls_2022_2026.csv).",
        "BCB: the bank's own contact centre. OVF: the overflow contractor that takes calls on the busiest days, which "
        "uses its own numeric codes.",
        "A firm has been reached once the centre has spoken with the respondent or with someone else at the firm, on a call "
        "the centre placed or on a call the firm made to the centre (records with reached = Y, and the contractor's codes "
        "that map to them). A report that the business has stopped trading does not count as reaching the firm.",
        "Calls the firm makes to the centre (direction in) are taken at the bank's inbound desk and carry the same codes.",
        "Code 71: the contractor's agent passed the call to a supervisor, whose record carries the rest of the same call and "
        "its outcome.",
        "BCB records carry the panel ID. The contractor's records carry the company number and region code from the "
        "contractor's own copy of the panel list.",
        "Owner: Contact Centre Manager (%s). Last revised 2 March 2022, when the centre moved to its current system."
        % W.PEOPLE["cc"]]:
        ws3.append([line])
    ws3.column_dimensions["A"].width = 125
    _save(wb, path, W.PEOPLE["cc"], "2022-03-02T08:20:00Z", "2022-03-02T15:05:00Z")


# ------------------------------------------------------------------------------------------------ data dictionary
FILES = [
    ("BCS_fieldwork_quality_standard_v3.pdf", "The fieldwork quality standard, version 3.0."),
    ("BCS_round_calendar.xlsx", "One row per round, March 2017 to December 2026."),
    ("bcs_panel_register.csv", "One row per firm recruited to the panel since 2014."),
    ("bcs_round_outcomes_2017_2026.csv", "One row per panel firm invited to a round, March 2017 to June 2026."),
    ("bcs_invitations_2026_09.csv", "One row per panel firm invited to the September 2026 round."),
    ("bcs_returns_2026_09.csv", "One row per return the fieldwork system logged for the September 2026 round up to the "
                                "main close; a firm can have more than one."),
    ("brevland_business_register_closures.csv", "One row per closure record in the register's extract of 25 September 2026."),
    ("business_register_extract_notes.pdf", "The register's notes on its closure extract."),
    ("contact_centre_calls_2022_2026.csv", "One row per call record in the contact centre's system, calls placed and calls "
                                           "received, March 2022 to the September 2026 main close."),
    ("contact_centre_outcome_codes.xlsx", "The contact centre's outcome codes, both schemes."),
    ("bcs_region_review_2023.csv", "Region codes and names before and after the 2023 regional review."),
    ("BCS_regional_balances_2026_06.xlsx", "Published results of the June 2026 round: net balances by region and industry "
                                           "group."),
    ("BCS_questionnaire_2026.pdf", "The 2026 questionnaire."),
    ("export_manifest.json", "What was extracted, from which system and when."),
]
FIELDS = [
    ("bcs_panel_register.csv", "panel_id", "Panel identifier."),
    ("bcs_panel_register.csv", "company_number", "Company number in the register's current form."),
    ("bcs_panel_register.csv", "region", "Region, R code (codes as set in the 2023 regional review)."),
    ("bcs_panel_register.csv", "nace_section", "NACE section of the firm's main activity as the Brevland Business Register "
                                               "held it when the firm was recruited (see the sheet Industry groups)."),
    ("bcs_panel_register.csv", "size_band", "Employment band when recruited."),
    ("bcs_panel_register.csv", "recruited", "Date the firm agreed to join the panel."),
    ("bcs_panel_register.csv", "recruitment_wave", "The recruitment the firm joined in: the 2014 panel, the yearly "
                                                   "refreshment for the March round, or the 2024 expansion."),
    ("bcs_panel_register.csv", "left_panel", "Date the firm left the panel; blank while it is on the panel."),
    ("bcs_panel_register.csv", "left_reason", "Why the firm left: closure registered, retired after six rounds without a "
                                              "return, or withdrew."),
    ("bcs_round_outcomes_2017_2026.csv", "round", "Survey round as year and quarter (2026Q2 is the June 2026 round)."),
    ("bcs_round_outcomes_2017_2026.csv", "panel_id", "Panel identifier."),
    ("bcs_round_outcomes_2017_2026.csv", "region", "The firm's region at the round: the bank's numeric codes up to the March "
                                                   "2023 round, R codes from the June 2023 round."),
    ("bcs_round_outcomes_2017_2026.csv", "mode", "How the firm's return came in: web, post or phone. Blank if no return "
                                                 "reached the bank by the round's data cut-off."),
    ("bcs_round_outcomes_2017_2026.csv", "received_at", "When the return was received into the survey database (UTC): web "
                                                        "returns when submitted, paper returns when the post room logged "
                                                        "them, telephone returns when the interview ended and, from 9 April "
                                                        "2024, when the contact centre's overnight transfer loaded them."),
    ("bcs_round_outcomes_2017_2026.csv", "recorded_at", "When the return's answers were recorded (UTC). Paper returns keyed "
                                                        "by the outside bureau from 1 February 2023 are recorded when "
                                                        "keyed."),
    ("bcs_invitations_2026_09.csv", "panel_id", "Panel identifier."),
    ("bcs_invitations_2026_09.csv", "company_number", "Company number in the register's current form."),
    ("bcs_invitations_2026_09.csv", "region", "Region, R code."),
    ("bcs_invitations_2026_09.csv", "region_name", "Region name."),
    ("bcs_invitations_2026_09.csv", "nace_section", "As in the panel register."),
    ("bcs_invitations_2026_09.csv", "size_band", "As in the panel register."),
    ("bcs_invitations_2026_09.csv", "invited_on", "Date the invitation went out."),
    ("bcs_invitations_2026_09.csv", "invitation_channel", "email, or post for firms that asked for paper."),
    ("bcs_returns_2026_09.csv", "receipt_no", "Fieldwork system receipt number."),
    ("bcs_returns_2026_09.csv", "panel_id", "Panel identifier."),
    ("bcs_returns_2026_09.csv", "mode", "web, post or phone."),
    ("bcs_returns_2026_09.csv", "received_at", "When the return reached the bank (UTC): web submission, end of the telephone "
                                               "interview, or post room log."),
    ("brevland_business_register_closures.csv", "company_number", "Company number (register's current form)."),
    ("brevland_business_register_closures.csv", "ceased_trading", "Date the company stopped trading."),
    ("brevland_business_register_closures.csv", "registered", "Date the closure was entered in the register."),
    ("brevland_business_register_closures.csv", "procedure", "Closure procedure (see the register's notes)."),
    ("contact_centre_calls_2022_2026.csv", "call_id", "Record number, assigned as records are loaded into the contact "
                                                      "centre's system."),
    ("contact_centre_calls_2022_2026.csv", "round", "Survey round (2026-09 is the September 2026 round)."),
    ("contact_centre_calls_2022_2026.csv", "team", "BCB: the bank's contact centre. OVF: the overflow contractor."),
    ("contact_centre_calls_2022_2026.csv", "direction", "out: a call the centre placed. in: a call a firm made to the centre."),
    ("contact_centre_calls_2022_2026.csv", "panel_id", "Panel identifier (BCB records)."),
    ("contact_centre_calls_2022_2026.csv", "company_number", "Company number from the contractor's list (OVF records)."),
    ("contact_centre_calls_2022_2026.csv", "region", "Region code from the contractor's list (OVF records)."),
    ("contact_centre_calls_2022_2026.csv", "agent", "Agent code."),
    ("contact_centre_calls_2022_2026.csv", "started_at", "When the call started (UTC)."),
    ("contact_centre_calls_2022_2026.csv", "duration_s", "Length of the call in seconds."),
    ("contact_centre_calls_2022_2026.csv", "outcome", "Outcome code (contact_centre_outcome_codes.xlsx)."),
    ("bcs_region_review_2023.csv", "old_code", "Region code before the review."),
    ("bcs_region_review_2023.csv", "old_name", "Region name before the review."),
    ("bcs_region_review_2023.csv", "new_code", "Region code from the review."),
    ("bcs_region_review_2023.csv", "new_name", "Region name from the review."),
    ("bcs_region_review_2023.csv", "effective_from", "Date the new codes took effect."),
]


def dictionary(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Files"
    _sheet(ws, ["file", "contents"], [list(x) for x in FILES], widths=[42, 110], wrap=True)
    ws2 = wb.create_sheet("Fields")
    _sheet(ws2, ["file", "field", "description"], [list(x) for x in FIELDS], widths=[38, 20, 110], wrap=True)
    ws3 = wb.create_sheet("Industry groups")
    rows = []
    for g, name in W.GROUPS:
        r21 = sorted(k for k, v in W.REV21_TO_GROUP.items() if v == g)
        rows.append([g, name, g, ", ".join(r21)])
    _sheet(ws3, ["survey group", "survey group name", "NACE Rev. 2 section", "NACE Rev. 2.1 sections"], rows,
           widths=[13, 42, 20, 22])
    ws3.append([])
    ws3.append(["The survey's nine industry groups are NACE Rev. 2 sections. The Brevland Business Register moved to NACE "
                "Rev. 2.1 on 1 January 2025, and firms recruited since then carry Rev. 2.1 sections."])
    ws4 = wb.create_sheet("Conventions")
    for line in [
        "Timestamps are UTC throughout. Brevland time is UTC+1, and UTC+2 from 01:00 UTC on the last Sunday in March to "
        "01:00 UTC on the last Sunday in October.",
        "Company numbers: six digits, a hyphen and a check digit, the register's form since 1 March 2021. Lists made before "
        "then can show the six digits alone.",
        "Regions: R01 to R12 since the 2023 regional review; the bank's older numeric codes before it "
        "(bcs_region_review_2023.csv).",
        "Owner: Survey Analyst (%s), Business Surveys Unit." % W.PEOPLE["analyst"]]:
        ws4.append([line])
    ws4.column_dimensions["A"].width = 125
    _save(wb, path, W.PEOPLE["analyst"], "2019-05-06T07:55:00Z", "2026-09-25T20:40:00Z")


# ------------------------------------------------------------------------------------------------ published balances (distractor)
def balances(path):
    r = records.rng("balances")
    wb = Workbook()
    ws = wb.active
    ws.title = "By region"
    ind = ["Output, past three months", "Output, next three months", "New orders, past three months",
           "Employment, next three months", "Selling prices, next three months", "Business situation"]
    base = np.array([4.0, 9.0, 2.0, 5.0, 21.0, 6.0])
    rows = []
    for code in W.REGS:
        v = base + r.normal(0, 6.5, size=len(ind))
        rows.append([code, W.NAME[code]] + [round(float(x), 1) for x in v])
    _sheet(ws, ["region", "region name"] + ind, rows, widths=[8, 18] + [16] * len(ind), wrap=False)
    ws2 = wb.create_sheet("By industry group")
    rows = []
    for g, name in W.GROUPS:
        v = base + r.normal(0, 8.0, size=len(ind))
        rows.append([g, name] + [round(float(x), 1) for x in v])
    _sheet(ws2, ["group", "group name"] + ind, rows, widths=[7, 36] + [16] * len(ind))
    ws3 = wb.create_sheet("Notes")
    for line in ["Business Conditions Survey, June 2026 round: published results, 30 July 2026.",
                 "Net balance: the percentage of firms reporting an increase less the percentage reporting a decrease, weighted "
                 "by employment. Business situation: good less poor.",
                 "Estimates use the returns received by the round's data cut-off."]:
        ws3.append([line])
    ws3.column_dimensions["A"].width = 110
    _save(wb, path, W.PEOPLE["analyst"], "2026-07-24T08:02:00Z", "2026-07-30T06:45:00Z")
