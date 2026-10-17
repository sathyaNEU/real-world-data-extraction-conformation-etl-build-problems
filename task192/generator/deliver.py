"""The two files the programme officer puts in front of the trustees."""
import world as W


def workbook(path, g):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="E8E8E4")

    def header(ws, row, names, widths):
        for j, h in enumerate(names):
            c = ws.cell(row=row, column=1 + j, value=h)
            c.font = Font(bold=True)
            c.fill = head_fill
            c.border = box
            c.alignment = Alignment(horizontal="left" if j == 0 else "center", wrap_text=True)
        for (col, wd) in widths:
            ws.column_dimensions[col].width = wd

    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = "Vigilancia Nacional: boletim do periodo de 1 de maio a 30 de outubro de 2026"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Indicador 2 do protocolo, ao mais proximo de cinquenta acoes. Registo extraido "
                "em 6 de novembro de 2026; o numero de acoes de confronto e o das acoes que o "
                "foram, nao o dos codigos atribuidos.")
    ws["A2"].font = Font(size=9, italic=True)
    rows = [("Acoes de confronto do periodo", "%d acoes" % g["total"]),
            ("Acoes registadas no periodo", "%d acoes" % g["actions_in_period"]),
            ("Acoes com relato do monitor de imprensa", "%d acoes" % g["press_only"]),
            ("Acoes so com relato de parceiro ou da linha", "%d acoes" % g["filed_by_hand"]),
            ("Codigo de confronto do classificador", "%d acoes" % g["press_marked"]),
            ("Codigo de confronto de quem entregou", "%d acoes" % g["hand_confrontational"]),
            ("Relatos na amostra de dupla codificacao", "%d relatos" % g["sample_rows"])]
    for i, (k, v) in enumerate(rows, 4):
        a = ws.cell(row=i, column=1, value=k)
        a.font = Font(bold=(i == 4))
        b = ws.cell(row=i, column=2, value=v)
        b.font = Font(bold=(i == 4))
        b.alignment = Alignment(horizontal="right")
    ws.cell(row=12, column=1, value=("As acoes com relato do monitor de imprensa levam o codigo do "
                                     "classificador (protocolo 2.6). As duas rondas de dupla "
                                     "codificacao do periodo fecharam em 20 de junho e em 10 de "
                                     "setembro de 2026 e cobrem os relatos de 1 de maio a 31 de "
                                     "agosto."))
    ws.cell(row=12, column=1).font = Font(size=9, italic=True)
    ws.cell(row=13, column=1, value="Preparado para o conselho de curadores de 19 de novembro de 2026.")
    ws.cell(row=13, column=1).font = Font(size=9, italic=True)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 20

    q = wb.create_sheet("Quinzenas")
    q["A1"] = "Acoes, territorio e rede de parceiros, por quinzena"
    q["A1"].font = Font(bold=True, size=12)
    header(q, 3, ["Quinzena", "De", "A", "Acoes registadas", "Municipios tocados",
                  "Quota da rede de parceiros, %"],
           [("A", 12), ("B", 13), ("C", 13), ("D", 17), ("E", 18), ("F", 26)])
    r = 4
    for label, v in sorted(g["a1"].items(), key=lambda kv: int(kv[0][1:])):
        for j, val in enumerate([label, v["start"], v["end"], v["actions"], v["municipalities"],
                                 v["partner_share"]]):
            c = q.cell(row=r, column=1 + j, value=val)
            c.border = box
            if j:
                c.alignment = Alignment(horizontal="center")
            if j in (3, 4):
                c.number_format = "#,##0"
            if j == 5:
                c.number_format = "0.0"
        r += 1
    q.cell(row=r + 1, column=1, value=("Uma acao e um tema numa localidade num dia, qualquer que "
                                       "seja o numero de relatos. A quota e a parte das acoes da "
                                       "quinzena que a rede de parceiros relatou. As quinzenas tem "
                                       "catorze dias e a ultima tem quinze."))
    q.cell(row=r + 1, column=1).font = Font(size=9, italic=True)
    q.freeze_panes = "A4"

    v = wb.create_sheet("Vias de acesso")
    v["A1"] = "Relatos por via de acesso no periodo"
    v["A1"].font = Font(bold=True, size=12)
    header(v, 3, ["Via", "Relatos recebidos", "Sem localidade fixada", "Minutos do inicio ao relato, mediana"],
           [("A", 30), ("B", 18), ("C", 22), ("D", 34)])
    r = 4
    name = dict(W.ROUTES)
    for code, _ in W.ROUTES:
        a2 = g["a2"][code]
        for j, val in enumerate(["%s (%s)" % (name[code], code), a2["received"], a2["unplaceable"],
                                 a2["median_minutes"]]):
            c = v.cell(row=r, column=1 + j, value=val)
            c.border = box
            if j:
                c.alignment = Alignment(horizontal="center")
            if j in (1, 2):
                c.number_format = "#,##0"
            if j == 3:
                c.number_format = "0.0"
        r += 1
    v.cell(row=r + 1, column=1, value=("Relatos recebidos em hora local no periodo, incluindo os "
                                       "que o gazeteiro nao fixou."))
    v.cell(row=r + 1, column=1).font = Font(size=9, italic=True)

    p = wb.create_sheet("Parceiros")
    p["A1"] = "Organizacoes parceiras: entregas e pagamento do periodo"
    p["A1"].font = Font(bold=True, size=12)
    header(p, 3, ["Codigo", "Organizacao", "Relatos entregues", "Entregues apos o fecho",
                  "Euros pagos"],
           [("A", 10), ("B", 46), ("C", 18), ("D", 22), ("E", 14)])
    r = 4
    for code, a3 in sorted(g["a3"].items()):
        for j, val in enumerate([code, a3["name"], a3["filed"], a3["after_close"], a3["euros"]]):
            c = p.cell(row=r, column=1 + j, value=val)
            c.border = box
            if j > 1:
                c.alignment = Alignment(horizontal="center")
            if j in (2, 3):
                c.number_format = "#,##0"
            if j == 4:
                c.number_format = "#,##0"
        r += 1
    p.cell(row=r + 1, column=1, value=("Relatos contados pelas acoes do periodo. Os entregues apos "
                                       "o fecho sao pagos com o periodo em que foram entregues "
                                       "(clausula 6.2)."))
    p.cell(row=r + 1, column=1).font = Font(size=9, italic=True)
    p.cell(row=r + 2, column=1, value=("A tabela do Observatorio Rural do Alentejo Central foi "
                                       "alterada pela adenda de 24 de julho, com efeito em 1 de "
                                       "agosto."))
    p.cell(row=r + 2, column=1).font = Font(size=9, italic=True)
    p.freeze_panes = "A4"
    v.freeze_panes = "A4"

    nt = wb.create_sheet("Notas")
    nt["A1"] = "Notas do boletim"
    nt["A1"].font = Font(bold=True, size=12)
    notes = [
        ("Registo", "Vigilancia Nacional, aberta em 1 de maio de 2026; extracao de 6 de novembro de 2026."),
        ("Indicador 2", "Acoes de confronto do protocolo 3.3, ao mais proximo de cinquenta acoes."),
        ("Acao", "Um tema de disputa numa localidade num dia (protocolo 2.1). Os relatos das tres "
                 "vias e das duas edicoes do grupo Atlantico contam uma vez."),
        ("Codigo da acao", "O do relato do monitor de imprensa quando existe, o de quem entregou nas "
                           "restantes acoes (protocolo 2.6)."),
        ("Quinzenas", "Catorze dias cada, contadas a partir de 1 de maio; os dias que sobram "
                      "entram na ultima, que tem quinze (protocolo 3.4)."),
        ("Carga de agosto", "O mes de agosto da rede de parceiros foi reposto numa segunda carga; "
                            "cada relato conta uma vez (registo_importacoes)."),
        ("Relogios", "O carimbo de recepcao do monitor de imprensa e o do portal de parceiros "
                     "vem em UTC; a hora de inicio de cada acao vem no relogio local do seu "
                     "distrito, e os Acores estao uma hora atras do continente."),
        ("Dupla codificacao", "Duas rondas, fechadas em 20 de junho e em 10 de setembro, sobre "
                              "os relatos aceites do monitor entrados desde o fecho anterior "
                              "(protocolo 4.5). A ronda de outono e do periodo seguinte."),
        ("Temas", "Dos catorze temas de disputa do protocolo (2.7), energia e redes e florestas e "
                  "fogos entraram no registo com a mobilizacao de setembro."),
        ("Contador de importacao", "O contador diario por via conta cada chegada, pelo que a carga "
                                   "reposta de agosto esta contada outra vez em 14 de setembro; "
                                   "cada relato conta uma vez."),
    ]
    for i, (k, vtxt) in enumerate(notes, 3):
        a = nt.cell(row=i, column=1, value=k)
        a.font = Font(bold=True)
        c = nt.cell(row=i, column=2, value=vtxt)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    nt.column_dimensions["A"].width = 20
    nt.column_dimensions["B"].width = 104
    wb.save(path)
    return path


def chart(path, g):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    labels = sorted(g["a1"], key=lambda k: int(k[1:]))
    total = [g["a1"][k]["actions"] for k in labels]
    conf = [g["a1"][k]["marked"] for k in labels]
    rest = [t - c for t, c in zip(total, conf)]
    fig, ax = plt.subplots(figsize=(8.6, 4.3), dpi=150)
    ax.bar(labels, rest, color="#9fb0c4", edgecolor="#ffffff", linewidth=1.1,
           label="acoes com outro codigo")
    ax.bar(labels, conf, bottom=rest, color="#8c2f2f", edgecolor="#ffffff", linewidth=1.1,
           label="acoes com codigo de confronto")
    for i, t in enumerate(total):
        ax.annotate("%d" % t, xy=(i, t), xytext=(0, 3), textcoords="offset points",
                    ha="center", fontsize=8.5)
    ax.set_ylabel("Acoes registadas")
    ax.set_title("Acoes registadas por quinzena, 1 de maio a 30 de outubro de 2026")
    ax.legend(frameon=False, fontsize=8.5, loc="upper left")
    ax.set_ylim(0, max(total) * 1.18)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(axis="x", labelsize=9)
    ax.grid(axis="y", color="#dcdfe4", linewidth=0.6)
    ax.set_axisbelow(True)
    fig.text(0.012, 0.012, "Vigilancia Nacional, registo extraido em 6 de novembro de 2026. As "
                           "quinzenas tem catorze dias e a ultima tem quinze.", fontsize=7.6,
             color="#55606b")
    fig.tight_layout(rect=(0, 0.035, 1, 1))
    fig.savefig(path, metadata={"Software": None, "Creation Time": None})
    plt.close(fig)
    return path


def bulletin(path, g, chart_png):
    """The bulletin page the association publishes for the period."""
    import os
    from reportlab import rl_config
    rl_config.invariant = 1
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_JUSTIFY
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["BodyText"], fontName="Times-Roman", fontSize=9.5,
                          leading=13, alignment=TA_JUSTIFY, spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="Times-Bold", fontSize=14, spaceAfter=6)
    big = ParagraphStyle("big", parent=ss["BodyText"], fontName="Times-Bold", fontSize=26,
                         leading=30, spaceAfter=2)
    small = ParagraphStyle("s", parent=body, fontSize=8, leading=10.5)
    cell = ParagraphStyle("c", fontName="Times-Roman", fontSize=8, leading=10)
    cellb = ParagraphStyle("cb", fontName="Times-Bold", fontSize=8, leading=10)
    tbl = TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.Color(0.45, 0.45, 0.45)),
                      ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.90, 0.90, 0.88)),
                      ("VALIGN", (0, 0), (-1, -1), "TOP"),
                      ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                      ("TOPPADDING", (0, 0), (-1, -1), 1.6), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6)])
    rows = [["Quinzena", "Acoes registadas", "Municipios tocados", "Quota da rede, %"]]
    for label in sorted(g["a1"], key=lambda k: int(k[1:])):
        v = g["a1"][label]
        rows.append([label, "%d" % v["actions"], "%d" % v["municipalities"],
                     "%.1f" % v["partner_share"]])
    d = SimpleDocTemplate(path, pagesize=A4, title="Boletim de acoes de confronto, 2026",
                          author="", subject="", creator="", leftMargin=18 * mm, rightMargin=18 * mm,
                          topMargin=16 * mm, bottomMargin=16 * mm)
    flow = [Paragraph("Vigilancia Nacional: acoes de confronto, 1 de maio a 30 de outubro de 2026", h1),
            Paragraph("Boletim do periodo, para o conselho de curadores de 19 de novembro de 2026. "
                      "Registo extraido em 6 de novembro de 2026.", small),
            Spacer(1, 6),
            Paragraph("%d" % g["total"], big),
            Paragraph("acoes de confronto no periodo, ao mais proximo de cinquenta, de %d acoes "
                      "registadas (indicador 2 do protocolo)." % g["actions_in_period"], body),
            Spacer(1, 4),
            Paragraph("Das acoes do periodo, %d tem relato do monitor de imprensa e levam o codigo "
                      "do classificador, e %d chegam so pela rede de parceiros ou pela linha de "
                      "participacao publica. O registo carrega codigo de confronto em %d acoes do "
                      "primeiro grupo e em %d do segundo. A mesa de codificacao codificou a mao %d "
                      "relatos do monitor nas duas rondas do periodo, fechadas em 20 de junho e em "
                      "10 de setembro."
                      % (g["press_only"], g["filed_by_hand"], g["press_marked"],
                         g["hand_confrontational"], g["sample_rows"]), body),
            Spacer(1, 6),
            Image(chart_png, width=168 * mm, height=84 * mm),
            Spacer(1, 6),
            Table([[Paragraph(str(c), cellb if i == 0 else cell) for c in r]
                   for i, r in enumerate(rows)],
                  colWidths=[26 * mm, 40 * mm, 42 * mm, 40 * mm], style=tbl, hAlign="LEFT"),
            Spacer(1, 4),
            Paragraph("Uma acao e um tema de disputa numa localidade num dia (protocolo 2.1), "
                      "qualquer que seja o numero de relatos. As quinzenas tem catorze dias e a "
                      "ultima tem quinze.", small)]
    d.build(flow)
    return path
