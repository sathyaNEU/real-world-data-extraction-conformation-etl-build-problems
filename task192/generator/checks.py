"""Every claim the build rests on, asserted against the cut bundle.

The three-rules block (R1, R2, R3a, R3b, R3c) runs on the register the bundle rebuilds and on the
clean truth behind it, so the client's three rules are executed rather than argued. The candidate
block executes every route to the indicator that the record leaves standing, names each one, and
asserts which cell of the stated granularity it lands in.
"""
import csv
import datetime as dt
import json
import math
import os
import statistics

import numpy as np

import golden as G
import world as W

N = 0

# the full candidate set: every route to the period's confrontational count that the shipped record
# leaves standing, with the family each one belongs to
RIGHT = ["rg_pooled", "matrix_inverse", "hybrid", "continuity", "rg_fb_month", "rg_fb_theme"]
WRONG = ["naive", "remove_fp_ppv", "remove_fp_fpr", "recall_only", "pv_pooled", "audit_prevalence",
         "hand_where_audited", "pv_fb_month", "pv_fb_theme"]


def ok(label, cond, detail=""):
    global N
    N += 1
    if not cond:
        raise AssertionError("%s failed %s" % (label, detail))


def rows_of(target, name):
    with open(os.path.join(target, name), encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def candidates(g):
    return {"naive": g["raw_rung0"], "remove_fp_ppv": g["raw_rung1"],
            "remove_fp_fpr": g["raw_rung1b"], "recall_only": g["raw_rung2"],
            "pv_pooled": g["raw_rung3"], "audit_prevalence": g["raw_rung4"],
            "hand_where_audited": g["raw_rung5"], "rg_pooled": g["raw_total"],
            "matrix_inverse": g["variant_matrix_inverse"], "hybrid": g["variant_hybrid"],
            "continuity": g["variant_continuity"], "rg_fb_month": g["variant_rg_fb_month"],
            "rg_fb_theme": g["variant_rg_fb_theme"], "pv_fb_month": g["variant_pv_fb_month"],
            "pv_fb_theme": g["variant_pv_fb_theme"]}


def two_by_two_p(a, b, c, d):
    """A chi-square p for a 2x2 table, without a dependency the shipped bundle does not need."""
    n = a + b + c + d
    if min(a + b, c + d, a + c, b + d) == 0:
        return 1.0
    chi = n * (a * d - b * c) ** 2 / float((a + b) * (c + d) * (a + c) * (b + d))
    return math.erfc(math.sqrt(chi / 2.0))


def dispersion(counts):
    m = statistics.mean(counts)
    return statistics.variance(counts) / m if m else 0.0


def dispersion_by_week(pairs):
    """Poisson dispersion against a weekday-specific mean: the register has a working week, so the
    unconditional dispersion of a daily count carries that cycle and only the conditional one says
    whether the noise itself is plausible."""
    groups = {}
    for day, n in pairs:
        groups.setdefault(day.weekday(), []).append(n)
    chi, df = 0.0, 0
    for wd, xs in groups.items():
        m = statistics.mean(xs)
        if m <= 0:
            continue
        chi += sum((x - m) ** 2 / m for x in xs)
        df += len(xs) - 1
    return chi / df if df else 0.0


def bootstrap(g, draws=20000, seed=7):
    """The sampling spread of the two families on this record, resampling the audit's four cells."""
    cells = (g["tp"], g["fn"], g["fp"], g["tn"])
    n = sum(cells)
    N_, O, H = g["press_only"], g["press_marked"], g["hand_confrontational"]
    rng = np.random.default_rng(seed)
    dr = rng.multinomial(n, np.array(cells, float) / n, size=draws)
    rg, pv = [], []
    for tp, fn, fp, tn in dr:
        if (tp + fn) == 0 or (fp + tn) == 0 or (tp + fp) == 0 or (fn + tn) == 0:
            continue
        se, fm = tp / float(tp + fn), fp / float(fp + tn)
        if se - fm <= 0:
            continue
        rg.append(H + (O - fm * N_) / (se - fm))
        pv.append(H + O * tp / float(tp + fp) + (N_ - O) * fn / float(fn + tn))
    rg, pv = np.array(rg), np.array(pv)
    gap = pv - rg
    return dict(sd_rg=float(rg.std()), sd_pv=float(pv.std()), sd_gap=float(gap.std()),
                mean_gap=float(gap.mean()), lo=float(np.percentile(rg, 2.5)),
                hi=float(np.percentile(rg, 97.5)))


def counts(acts, sample):
    inper = [a for a in acts if W.PERIOD_FROM <= a["day"] <= W.PERIOD_TO]
    press = [a for a in inper if a["press_only"]]
    other = [a for a in inper if not a["press_only"]]
    H = sum(1 for a in other if a["coded_conf"])
    N_ = len(press)
    O = sum(1 for a in press if a["coded_conf"])
    cells = G.matrix(sample)
    raw, se, fm = G.rogan_gladen(N_, O, cells)
    tp, fn, fp, tn = cells
    return dict(n=len(inper), N=N_, H=H, O=O, g=H + raw, r0=H + O,
                r1=H + O * tp / float(tp + fp), r2=H + O / se,
                r3=H + G.predictive(N_, O, cells), se=se, fm=fm)


def run(target, g=None):
    global N
    N = 0
    g = g or G.compute(target)
    b = G.load(target)
    acts = G.actions(b["reports"])
    files = sorted(os.listdir(target))
    from openpyxl import load_workbook
    _wb = load_workbook(os.path.join(target, "dicionario_dados.xlsx"), read_only=True)
    conv = " ".join(str(c.value) for row in _wb["convencoes"].iter_rows() for c in row
                    if c.value is not None)
    _wb.close()
    cand = candidates(g)
    cell = {k: G.grain(v) for k, v in cand.items()}
    gold_cell = G.grain(g["raw_total"])

    # ---------------------------------------------------------------- input gates
    ok("gate.files", len(files) >= 10, "%d files" % len(files))
    ok("gate.formats", len({os.path.splitext(f)[1] for f in files}) >= 3)
    big = max(sum(1 for _ in open(os.path.join(target, f), encoding="utf-8")) - 1
              for f in files if f.endswith(".csv"))
    ok("gate.bulk", big >= 25000, "%d rows" % big)
    mpath = os.path.join(os.path.dirname(target), "metadata.json")
    if not os.path.exists(mpath):
        mpath = os.path.join(os.path.dirname(os.path.abspath(__file__)), os.pardir,
                             "metadata.json")
    meta = json.load(open(mpath))
    ok("gate.distractors", len(meta["distractor_files"]) >= 2)
    for dd in meta["distractor_files"]:
        ok("gate.distractor.present", dd["file"] in files, dd["file"])
    text = ""
    for f in files:
        if f.endswith((".csv", ".json")):
            text += open(os.path.join(target, f), encoding="utf-8", errors="replace").read(600000)
    ok("gate.no_answer_key", not any(k in text.lower() for k in
                                     ("is_correct", "should_exclude", "verdadeiro_codigo",
                                      "sensibilidade", "especificidade", "taxa_de_erro")))

    # ---------------------------------------------------------------- the call reproduces
    ok("call.actions", g["actions_in_period"] == 4627, g["actions_in_period"])
    ok("call.press_reported", g["press_only"] == 4200, g["press_only"])
    ok("call.other_actions", g["filed_by_hand"] == 427, g["filed_by_hand"])
    ok("call.hand_conf", g["hand_confrontational"] == 30, g["hand_confrontational"])
    ok("call.press_marked", g["press_marked"] == 302, g["press_marked"])
    ok("call.sample", g["sample_rows"] == 1000, g["sample_rows"])
    ok("call.matrix", (g["tp"], g["fn"], g["fp"], g["tn"]) == (74, 76, 21, 829),
       str((g["tp"], g["fn"], g["fp"], g["tn"])))
    ok("call.total", g["total"] == 450, g["total"])
    ok("call.raw", abs(g["raw_total"] - 453.0126) < 0.01, "%.4f" % g["raw_total"])
    ok("call.grain", g["grain"] == 50 and G.GRAIN == 50)
    ok("call.recall", abs(g["recall"] - 0.4933) < 0.0005, g["recall"])
    ok("call.false_mark", abs(g["false_mark"] - 0.0247) < 0.0005, g["false_mark"])
    ok("call.parts", g["press_only"] + g["filed_by_hand"] == g["actions_in_period"])

    # ---------------------------------------------------------------- the candidate set
    for k in RIGHT + WRONG:
        ok("cand.present.%s" % k, k in cand and cand[k] == cand[k])
    for k in RIGHT:
        ok("cand.converge.%s" % k, cell[k] == gold_cell,
           "%s %.3f -> %d, golden cell %d" % (k, cand[k], cell[k], gold_cell))
    for k in WRONG:
        ok("cand.separated.%s" % k, cell[k] != gold_cell,
           "%s %.3f lands in the golden's cell %d" % (k, cand[k], gold_cell))
        ok("cand.six_percent.%s" % k,
           abs(cand[k] - g["raw_total"]) >= 0.06 * g["raw_total"],
           "%s %.3f is %.2f per cent from %.3f"
           % (k, cand[k], 100 * abs(cand[k] - g["raw_total"]) / g["raw_total"], g["raw_total"]))
    ok("cand.identities", abs(cand["matrix_inverse"] - cand["rg_pooled"]) < 1e-9
       and abs(cand["hybrid"] - cand["rg_pooled"]) < 1e-9,
       "%.6f %.6f" % (cand["matrix_inverse"], cand["hybrid"]))
    ok("cand.nearest_wrong",
       min(abs(cand[k] - g["raw_total"]) for k in WRONG) >= 0.06 * g["raw_total"],
       "%.2f" % min(abs(cand[k] - g["raw_total"]) for k in WRONG))
    ok("cand.no_empty_cell", min(g["tp"], g["fn"], g["fp"], g["tn"]) > 0,
       "a continuity adjustment is for an empty cell and no cell is empty")
    ok("cand.distinct_cells", len({cell[k] for k in WRONG}) >= 4,
       str(sorted({cell[k] for k in WRONG})))

    # ---------------------------------------------------------------- the stratified families die
    for name, d in sorted(g["dead"].items()):
        ok("dead.%s" % name, d["value"] is None, "%s is computable" % name)
        ok("dead.%s.reason" % name, len(d["broken"]) > 0, name)
    ok("dead.month_reason", set(g["dead"]["rg_by_month"]["broken"]) == {"9", "10"},
       str(g["dead"]["rg_by_month"]["broken"]))
    ok("dead.theme_reason", set(g["dead"]["rg_by_theme"]["broken"]) >= set(W.LATE_THEMES),
       str(g["dead"]["rg_by_theme"]["broken"]))
    ok("dead.late_themes_after_surge",
       all(min(a["day"] for a in acts if a["theme"] == t) >= W.SURGE_FROM for t in W.LATE_THEMES),
       str(W.LATE_THEMES))
    ok("dead.late_themes_present",
       all(sum(1 for a in acts if a["theme"] == t) > 40 for t in W.LATE_THEMES))
    ok("dead.audit_window", g["audit_covers"] == [5, 6, 7, 8], str(g["audit_covers"]))
    ok("dead.population_wider",
       sorted({a["day"].month for a in acts
               if W.PERIOD_FROM <= a["day"] <= W.PERIOD_TO}) == [5, 6, 7, 8, 9, 10])
    ok("dead.rounds_disjoint",
       all(W.SAMPLE_ROUNDS[i][3] < W.SAMPLE_ROUNDS[i + 1][2]
           for i in range(len(W.SAMPLE_ROUNDS) - 1)))

    # ------------------------------------------------- the data forces the method, and shows it
    audited = [a for a in acts if W.PERIOD_FROM <= a["day"] <= W.PERIOD_TO and a["press_only"]
               and a["day"] <= W.SAMPLE_ROUNDS[-1][3]]
    rest = [a for a in acts if W.PERIOD_FROM <= a["day"] <= W.PERIOD_TO and a["press_only"]
            and a["day"] > W.SAMPLE_ROUNDS[-1][3]]
    ma = sum(1 for a in audited if a["coded_conf"])
    mr = sum(1 for a in rest if a["coded_conf"])
    p_marked = two_by_two_p(ma, len(audited) - ma, mr, len(rest) - mr)
    ok("force.marked_rate_moves", p_marked < 0.001,
       "the marked rate is %.4f inside the audited window and %.4f outside it, p %.2e"
       % (ma / float(len(audited)), mr / float(len(rest)), p_marked))
    ok("force.marked_rate_direction", ma / float(len(audited)) > 1.5 * mr / float(len(rest)),
       "%.4f %.4f" % (ma / float(len(audited)), mr / float(len(rest))))
    ok("force.audit_prevalence_above_record",
       g["audit_prevalence"] >= 1.4 * g["record_prevalence"],
       "audit %.4f record %.4f" % (g["audit_prevalence"], g["record_prevalence"]))
    r1, r2 = g["audit_rounds"]["R1"], g["audit_rounds"]["R2"]
    p_prev = two_by_two_p(r1["cells"][0] + r1["cells"][1],
                          r1["cells"][2] + r1["cells"][3],
                          r2["cells"][0] + r2["cells"][1],
                          r2["cells"][2] + r2["cells"][3])
    p_recall = two_by_two_p(r1["cells"][0], r1["cells"][1], r2["cells"][0], r2["cells"][1])
    p_fm = two_by_two_p(r1["cells"][2], r1["cells"][3], r2["cells"][2], r2["cells"][3])
    ok("force.rates_homogeneous", p_recall > 0.15 and p_fm > 0.15,
       "recall p %.3f, false mark p %.3f" % (p_recall, p_fm))
    ok("force.prevalence_falls_in_audit", r1["prevalence"] > r2["prevalence"],
       "%.4f %.4f" % (r1["prevalence"], r2["prevalence"]))
    months = sorted(g["audit_by_month"])
    prevs = [g["audit_by_month"][m][1] / float(g["audit_by_month"][m][0]) for m in months]
    ok("force.audit_prevalence_trend", prevs[0] > prevs[-1],
       str([round(x, 4) for x in prevs]))
    ok("force.model_fact_shipped",
       "treinado_em" in open(os.path.join(target, "config_monitor.json"), encoding="utf-8").read())

    # ------------------------------------------------- the separation is wider than the noise
    boot = bootstrap(g)
    g["boot"] = boot
    ok("noise.gap_beats_sd", boot["mean_gap"] >= 2.0 * boot["sd_gap"],
       "mean gap %.1f, sd of the gap %.1f, ratio %.2f"
       % (boot["mean_gap"], boot["sd_gap"], boot["mean_gap"] / boot["sd_gap"]))
    ok("noise.gap_observed", (cand["pv_pooled"] - cand["rg_pooled"]) >= 2.0 * boot["sd_gap"],
       "observed gap %.1f against sd %.1f"
       % (cand["pv_pooled"] - cand["rg_pooled"], boot["sd_gap"]))
    ok("noise.grain_not_finer_than_sd", G.GRAIN <= boot["sd_rg"],
       "grain %d, sampling sd %.1f" % (G.GRAIN, boot["sd_rg"]))
    daily = {}
    for a in acts:
        if W.PERIOD_FROM <= a["day"] <= W.SURGE_FROM - dt.timedelta(days=1):
            daily[a["day"]] = daily.get(a["day"], 0) + 1
    disp = dispersion_by_week(sorted(daily.items()))
    ok("noise.dispersion_plausible", 0.8 <= disp <= 1.35,
       "chi-square per degree of freedom %.3f on daily action counts before the surge, against a "
       "weekday mean" % disp)
    ok("noise.weekly_cycle_present", dispersion(list(daily.values())) > 1.2,
       "unconditional dispersion %.3f" % dispersion(list(daily.values())))
    conf_daily = {}
    for a in acts:
        if W.PERIOD_FROM <= a["day"] <= W.PERIOD_TO:
            conf_daily[a["day"]] = conf_daily.get(a["day"], 0) + (1 if a["coded_conf"] else 0)
    dm = dispersion_by_week(sorted(conf_daily.items()))
    ok("noise.dispersion_marked", 0.75 <= dm <= 1.4, "%.3f" % dm)

    # ---------------------------------------------------------------- the ladder and the bins
    order = [cand["remove_fp_fpr"], cand["remove_fp_ppv"], cand["naive"],
             cand["hand_where_audited"], cand["rg_pooled"], cand["pv_pooled"],
             cand["recall_only"], cand["audit_prevalence"]]
    ok("ladder.order", all(order[i] < order[i + 1] for i in range(len(order) - 1)),
       str([round(x, 1) for x in order]))
    ok("ladder.bracketed", cand["naive"] < cand["rg_pooled"] < cand["pv_pooled"])
    ok("ladder.decisive_move_largest",
       (cand["pv_pooled"] - cand["rg_pooled"]) > (cand["rg_pooled"] - cand["hand_where_audited"]))
    half = G.GRAIN / 2.0
    edge = min(abs(g["raw_total"] - (gold_cell - half)), abs((gold_cell + half) - g["raw_total"]))
    ok("bin.total", edge >= 12.0, "the call sits %.2f from its nearest bin edge" % edge)
    ok("bin.total_mid", abs(g["raw_total"] - gold_cell) <= 0.55 * half,
       "%.3f against the bin centre %d" % (g["raw_total"], gold_cell))
    for k in WRONG:
        e = min(abs(cand[k] - (cell[k] - half)), abs((cell[k] + half) - cand[k]))
        ok("bin.cand.%s" % k, e >= 1.5, "%s %.3f sits %.2f from its bin edge" % (k, cand[k], e))
    for code, v in sorted(g["a3"].items()):
        ok("bin.euros.%s" % code, 0.5 - abs(v["raw_euros"] - round(v["raw_euros"])) >= 0.25,
           "%.4f" % v["raw_euros"])
    # the fortnight's partner column is a count of actions, so it has no rounding and no edge. the
    # measurement that named the wrong share is kept as a record rather than as a gate: a share is
    # printed to one decimal, so its edge is at x.x5 and the distance to that edge is half a tenth
    # LESS the distance to the nearest printed value, not that distance itself.
    two_step, edges = [], {}
    for label, a, bb in W.fortnights():
        inside = [x for x in acts if a <= x["day"] <= bb]
        par = [x for x in inside if "PAR" in x["routes"]]
        ok("bin.partner_count.%s" % label,
           float(len(par)).is_integer() and 0 < len(par) < len(inside),
           "%s %d of %d actions, a whole count with nothing to round"
           % (label, len(par), len(inside)))
        share = 100.0 * len(par) / len(inside)
        tenths = share * 10.0
        edges[label] = round((0.5 - abs(tenths - round(tenths))) / 10.0, 4)
        if round(round(share, 2), 1) != round(share, 1):
            two_step.append("%s %.6f (%.1f against %.1f)"
                            % (label, share, round(round(share, 2), 1), round(share, 1)))
    g["share_edges"] = edges
    g["share_two_step"] = two_step
    # no graded figure is a percentage, so axis 10 carries no percentage at all; the bulletin's own
    # rounding rule for a percentage is filed anyway, which closes the axis for anything later
    ok("bin.share_rounding_filed",
       "arredondada uma vez" in conv and "nao truncada" in conv,
       "the percentage rule is in the dictionary's convencoes sheet")
    here = os.path.dirname(os.path.abspath(__file__))
    body = open(os.path.join(here, os.pardir, "submission.md"),
                encoding="utf-8").read().split("## 4. Deliverable Answers", 1)[-1]
    ok("bin.no_graded_percentage", "per cent" not in body and "%" not in body,
       "the deliverable answers state no percentage, so no graded figure is one")
    ok("bin.prompt_rounds_no_share",
       "share" not in open(os.path.join(here, os.pardir, "prompt.txt"), encoding="utf-8").read(),
       "the prompt asks for no share")
    for route, v in g["a2"].items():
        ok("bin.median.%s" % route, abs(v["median_minutes"] * 10
                                        - round(v["median_minutes"] * 10)) < 1e-9,
           "%.4f" % v["median_minutes"])
        ok("bin.median_pair.%s" % route, v["median_pair_equal"], route)

    # ---------------------------------------------------------------- R1, raw material
    heads = []
    for f in files:
        if f.endswith(".csv"):
            with open(os.path.join(target, f), encoding="utf-8") as fh:
                heads += next(csv.reader(fh))
    ok("R1.raw_rows", "codigo_monitor" in heads and "codigo_manual" in heads)
    ok("R1.no_rate_column", not any(h.lower().startswith(("taxa", "sensib", "especif", "erro"))
                                    for h in heads), str(heads))
    ok("R1.no_truth_column", not any("verdad" in h.lower() or "real" in h.lower() for h in heads))
    base = counts(acts, b["sample"])
    ok("R1.register_from_rows", base["n"] == 4627 and G.grain(base["g"]) == 450)

    # ---------------------------------------------------------------- R2, mess neutral
    def call_of(reports):
        return counts(G.actions(reports), b["sample"])

    ok("R2.base", G.grain(call_of(b["reports"])["g"]) == 450)
    one_load = [r for r in b["reports"] if r.get("load") != "C20260802"]
    ok("R2.reimport", G.grain(call_of(one_load)["g"]) == 450 and call_of(one_load)["n"] == 4627)
    shifted = [dict(r, recv=r["recv"] + dt.timedelta(hours=1)) for r in b["reports"]]
    ok("R2.clocks", G.grain(call_of(shifted)["g"]) == 450)
    cur = {r["codigo_localidade"]: r["nome"] for r in rows_of(target, "gazeteiro_localidades.csv")
           if r["estado_nome"] == "atual"}
    renamed = [dict(r, place=cur.get(r["loc"], r["place"])) for r in b["reports"]]
    ok("R2.former_names", G.grain(call_of(renamed)["g"]) == 450)
    items = rows_of(target, "registo_itens_imprensa.csv")
    per_report = {}
    for r in items:
        if r["estado"] == "aceite":
            per_report.setdefault(r["id_relato"], set()).add(r["codigo_acao"])
    ok("R2.one_code_per_report", all(len(v) == 1 for v in per_report.values()))
    ok("R2.items_above_reports", len([r for r in items if r["estado"] == "aceite"]) > len(per_report))
    ids = {r["ident"] for r in b["reports"]}
    ok("R2.exceptions_disjoint", not ({e["ident"] for e in b["exceptions"]} & ids))

    # ---------------------------------------------------------------- R3, the three rules executed
    ok("R3a.analyst_test", G.grain(base["r3"]) == 600 and G.grain(base["g"]) == 450
       and G.grain(base["r3"]) != G.grain(base["g"]),
       "the default path files %d, the golden is %d" % (G.grain(base["r3"]), G.grain(base["g"])))
    ok("R3a.codes_as_filed_wrong", G.grain(base["r0"]) == 350)
    ok("R3b.decisive_only", G.grain(base["r2"]) == 650 and G.grain(base["g"]) == 450)
    ok("R3b.each_rung_wrong", 450 not in {250, 350, 400, 600, 650})
    with open(os.path.join(target, "registo_itens_imprensa.csv"), encoding="utf-8") as fh:
        lines = fh.read().splitlines()[1:]
    blocks = {}
    for i in range(0, len(lines) - 20, 20):
        key = "|".join(x.split(",", 1)[1] for x in lines[i:i + 20])
        blocks[key] = blocks.get(key, 0) + 1
    ok("R3c.no_copied_block", max(blocks.values()) == 1)
    for word in ("sensib", "especif", "recall", "precis", "falso positivo", "falso negativo",
                 "corrig", "taxa de erro", "representativ", "extrapol"):
        ok("R3c.no_method_word.%s" % word[:9], word not in text.lower(), word)
    pdfs = "".join(open(os.path.join(target, f), "rb").read().decode("latin-1")
                   for f in files if f.endswith(".pdf"))
    for word in ("sensibilidade", "especificidade", "falso positivo", "falso negativo",
                 "representativ", "453", "450 acoes"):
        ok("R3c.no_method_in_pdf.%s" % word[:9], word not in pdfs, word)

    # ---------------------------------------------------------------- the semantic fork table
    ok("SF1.action_grain", g["actions_in_period"] < len(b["reports"]))
    ok("SF2.codes_not_actions", G.grain(g["raw_rung0"]) != g["total"])
    ok("SF3.codes", set(W.CONFRONTATIONAL) == {"CF", "RF", "DP"})
    ok("SF4.window", all(W.PERIOD_FROM <= a["day"] <= W.PERIOD_TO
                         for a in acts if W.PERIOD_FROM <= a["day"] <= W.PERIOD_TO))
    ok("SF5.code_kept", all((a["code"] == a["press_code"]) if a["press_code"] else True
                            for a in acts))
    ok("SF6.sample_scope", {r["id_relato"][:3] for r in b["sample"]} == {"IMP"})
    ok("SF7.action_date_fixed", all(r["start"].date() == r["day"] for r in b["reports"]))
    ok("SF8.place_resolution", all(r["loc"] for r in b["reports"]))
    ok("SF9.reimport_neutral", G.grain(call_of(one_load)["g"]) == 450)
    ok("SF10.twin_in_items", any(r["orgao"] == W.TWIN_FEED[1] for r in items))
    ok("SF11.rounding", abs(g["raw_total"] - gold_cell) <= 0.55 * half)
    ok("SF12.ties", len({250, 350, 400, 450, 600, 650}) == 6)
    ok("SF13.accepted_only", all(r["estado"] in ("aceite", "rejeitado") for r in items))
    ok("SF14.extract", max(r["day"] for r in b["reports"]) <= W.EXTRACT)
    ok("SF15.outside_period", any(a["day"] > W.PERIOD_TO for a in acts))
    ok("SF16.hand_coding_present", all(r["codigo_manual"] for r in b["sample"]))
    ok("SF17.unplaced_not_actions", len(b["exceptions"]) > 0)
    ok("SF18.no_carry_in", min(r["day"] for r in b["reports"]) >= W.OPEN_DAY)
    ok("SF19.fortnight_days", all((bb - a).days + 1 == W.FORTNIGHT_DAYS
                                  for _, a, bb in W.fortnights()[:-1]))
    ok("SF20.last_fortnight", (W.fortnights()[-1][2] - W.fortnights()[-1][1]).days + 1
       == W.FORTNIGHT_DAYS + 1)
    ok("SF21.delivery_codes_mapped",
       W.MERGED[1] in {r["codigo_parceiro"] for r in rows_of(target,
                                                             "relatos_rede_parceiros.csv")})
    ok("SF22.azores_split", {"Acores", "Madeira"} <=
       {r["distrito"] for r in rows_of(target, "gazeteiro_localidades.csv")})
    cfg = json.load(open(os.path.join(target, "config_monitor.json"), encoding="utf-8"))
    ok("SF23.offsets_shipped", "fusos_horarios" in cfg
       and "acores" in cfg["fusos_horarios"] and "25 de outubro" in str(cfg["fusos_horarios"]))
    ok("SF24.azores_actions_present",
       any(a["district"] == W.AZORES_DISTRICT for a in acts))
    ok("SF25.route_reported_filed", "ainda que outras vias" in conv,
       "the convention that a route reported an action even when other routes did is shipped")
    ok("SF25.route_reported_bites",
       sum(1 for a in acts if W.PERIOD_FROM <= a["day"] <= W.PERIOD_TO
           and a["routes"] == {"PAR"}) < sum(1 for a in acts
                                             if W.PERIOD_FROM <= a["day"] <= W.PERIOD_TO
                                             and "PAR" in a["routes"]),
       "the two readings of a reported action give different counts")

    # ---------------------------------------------------------------- realism
    by_fn = [v["actions"] for v in g["a1"].values()]
    ok("real.no_identical_periods", max(by_fn.count(x) for x in by_fn) <= 2, str(by_fn))
    ok("real.no_round_call", g["raw_total"] % 1 != 0)
    for route, a2 in g["a2"].items():
        ok("real.median_not_round.%s" % route, a2["median_minutes"] % 10 != 0,
           "%s %s" % (route, a2["median_minutes"]))
    for p, a3 in g["a3"].items():
        ok("real.euros_not_round.%s" % p, a3["euros"] % 100 != 0, "%s %s" % (p, a3["euros"]))
    for name in W.PEOPLE.values():
        hit = [f for f in files if name.split()[0].encode() in open(os.path.join(target, f), "rb").read()]
        ok("real.persona_only_where_filed.%s" % name.split()[0],
           all(f.endswith((".pdf", ".xlsx")) for f in hit), str(hit))
    ok("real.coders_are_initials",
       {r["codificador"] for r in b["sample"]} <= {c for c, _ in W.CODERS})
    ok("real.coder_loads_differ", len({sum(1 for r in b["sample"] if r["codificador"] == c)
                                       for c, _ in W.CODERS}) == len(W.CODERS))
    load = sorted(sum(1 for r in b["sample"] if r["codificador"] == c) for c, _ in W.CODERS)
    ok("real.coder_spread", load[-1] >= 1.5 * load[0], str(load))
    ok("real.sample_two_rounds", {r["ronda"] for r in b["sample"]} == {"R1", "R2"})
    ok("real.adjudication_present",
       0.04 <= sum(1 for r in b["sample"] if r["revisto_por"]) / float(len(b["sample"])) <= 0.14)
    ok("real.adjudication_decided",
       {r["decisao_da_mesa"] for r in b["sample"] if r["revisto_por"]} == {"mantido", "alterado"})
    ok("real.notes_present", 10 <= sum(1 for r in b["sample"] if r["nota"]) <= 90)
    ok("real.coding_after_action",
       all(r["data_codificacao"] > r["data_ocorrencia"] for r in b["sample"]))
    ok("real.coding_before_extract",
       max(r["data_codificacao"] for r in b["sample"]) <= W.EXTRACT.isoformat())
    cal = {r["data_codificacao"] for r in b["sample"]}
    ok("real.coding_calendar_uneven", len(cal) >= 30 and max(
        sum(1 for r in b["sample"] if r["data_codificacao"] == dd) for dd in cal) >= 2 * min(
        sum(1 for r in b["sample"] if r["data_codificacao"] == dd) for dd in cal))
    mins = [int(r["inicio_acao"][14:16]) for r in items if r["inicio_acao"]]
    ok("real.minutes_not_all_round",
       0.6 <= sum(1 for m in mins if m % 5 == 0) / float(len(mins)) <= 0.92,
       "%.3f on multiples of five" % (sum(1 for m in mins if m % 5 == 0) / float(len(mins))))
    hours = {int(r["inicio_acao"][11:13]) for r in items if r["inicio_acao"]}
    ok("real.midday_present", {12, 13} <= hours, str(sorted(hours)))
    weekday, weekend = [], []
    for day, n in daily.items():
        (weekend if day.weekday() >= 5 else weekday).append(n)
    ok("real.weekend_effect", statistics.mean(weekend) <= 0.8 * statistics.mean(weekday),
       "%.1f against %.1f" % (statistics.mean(weekend), statistics.mean(weekday)))
    par_hours = [int(r["carimbo_entrega_utc"][11:13])
                 for r in rows_of(target, "relatos_rede_parceiros.csv")]
    peak = max(par_hours.count(h) for h in set(par_hours))
    ok("real.partner_hours_diurnal", peak <= 0.30 * len(par_hours)
       and len({h for h in set(par_hours) if par_hours.count(h) > len(par_hours) / 48.0}) >= 8,
       "busiest hour holds %.3f of the deliveries" % (peak / float(len(par_hours))))
    lags = [r["lag"] for r in b["reports"] if r["route"] == "LIN" and r["lag"] is not None]
    ok("real.line_lag_spread", max(lags) / min(max(lags[0], 1), max(1, min(lags))) >= 20,
       "%d to %d minutes" % (min(lags), max(lags)))
    agree = sum(1 for r in b["sample"] if r["codigo_monitor"] == r["codigo_manual"])
    ok("real.classifier_is_an_instrument", agree / float(len(b["sample"])) >= 0.45,
       "%.3f exact code agreement" % (agree / float(len(b["sample"]))))
    near = far = 0
    pos = {c: i for i, (c, _, _) in enumerate(W.ACTION_CODES)}
    for r in b["sample"]:
        if r["codigo_monitor"] == r["codigo_manual"]:
            continue
        if abs(pos[r["codigo_monitor"]] - pos[r["codigo_manual"]]) <= 2:
            near += 1
        else:
            far += 1
    ok("real.confusion_is_local", near > far, "near %d far %d" % (near, far))
    ok("real.extract_coherent", all(r["recv"].date() <= W.EXTRACT for r in b["reports"]))
    man = {r["ficheiro"]: r for r in rows_of(target, "manifesto_extracao.csv")}
    ok("real.manifest_covers", set(files) - set(man) == {"manifesto_extracao.csv"},
       str(set(files) - set(man)))
    ok("real.manifest_extract_date",
       all(r["extraido_em"] == W.EXTRACT.isoformat() for r in man.values()))
    ok("real.item_rows_match", int(man["registo_itens_imprensa.csv"]["linhas"]) == big)

    # ---------------------------------------------------------------- the ask layer
    ok("ask.a1_fortnights", len(g["a1"]) == 13)
    ok("ask.a1_last_fortnight", (W.fortnights()[-1][2] - W.fortnights()[-1][1]).days == 14)
    ok("ask.a2_routes", len(g["a2"]) == 3 and all(v["unplaceable"] > 0 for v in g["a2"].values()))
    ok("ask.a3_partners", len(g["a3"]) == 7 and all(v["filed"] > 0 for v in g["a3"].values()))
    par_rows = rows_of(target, "relatos_rede_parceiros.csv")
    ok("cracker.a2_received", len(par_rows) != g["a2"]["PAR"]["received"],
       "%d %d" % (len(par_rows), g["a2"]["PAR"]["received"]))
    counter = {}
    for r in b["totals"]:
        counter[r["via"]] = counter.get(r["via"], 0) + int(r["relatos_recebidos"])
    ok("cracker.a2_counter_overcounts", counter["PAR"] > g["a2"]["PAR"]["received"],
       "counter %d against %d" % (counter["PAR"], g["a2"]["PAR"]["received"]))
    for route in ("IMP", "LIN"):
        inside = sum(int(r["relatos_recebidos"]) for r in b["totals"] if r["via"] == route
                     and W.PERIOD_FROM <= G.pdate(r["data"]) <= W.PERIOD_TO)
        ok("ask.a2_counter_reconciles.%s" % route, inside == g["a2"][route]["received"],
           "%s counter %d against %d" % (route, inside, g["a2"][route]["received"]))
    ok("cracker.a2_unplaceable", g["a2"]["IMP"]["unplaceable"] > 0)
    naive_med = statistics.median(sorted(
        (r["recv_utc"] - r["start"]).total_seconds() / 60.0
        for r in b["reports"] if r["route"] == "IMP"
        and W.PERIOD_FROM <= r["recv"].date() <= W.PERIOD_TO))
    ok("cracker.a2_median", abs(naive_med - g["a2"]["IMP"]["median_minutes"]) > 5,
       "%.1f %s" % (naive_med, g["a2"]["IMP"]["median_minutes"]))
    seen = {}
    for r in b["reports"]:
        if r["route"] == "PAR":
            seen.setdefault(r["ident"], r)
    flat = sum(dict((p[0], p[3]) for p in W.PARTNERS)[W.AMENDED[0]] for r in seen.values()
               if r["src"] == W.AMENDED[0] and W.PERIOD_FROM <= r["day"] <= W.PERIOD_TO
               and r["recv"].date() <= W.PERIOD_TO)
    ok("cracker.a3_amendment", round(flat) != g["a3"][W.AMENDED[0]]["euros"],
       "%d %d" % (round(flat), g["a3"][W.AMENDED[0]]["euros"]))
    one_code = sum(1 for r in seen.values() if r["src"] == W.MERGED[0]
                   and W.PERIOD_FROM <= r["day"] <= W.PERIOD_TO)
    ok("cracker.a3_merged", one_code != g["a3"][W.MERGED[0]]["filed"],
       "%d %d" % (one_code, g["a3"][W.MERGED[0]]["filed"]))
    all_rows = sum(1 for r in seen.values())
    ok("cracker.a3_window", all_rows != sum(v["filed"] for v in g["a3"].values()),
       "%d %d" % (all_rows, sum(v["filed"] for v in g["a3"].values())))
    ok("cracker.a1_scope", any(a["day"] > W.PERIOD_TO for a in acts))

    # ---------------------------------------------------------------- the separation line
    ok("sep.devices_outside_call", all(e["ident"] not in ids for e in b["exceptions"]))
    ok("sep.totals_carry_no_code",
       set(rows_of(target, "totais_por_via.csv")[0]) == {"data", "via", "relatos_recebidos"})
    ok("sep.sample_carries_no_total",
       "total" not in ",".join(rows_of(target, "amostra_dupla_codificacao.csv")[0]).lower())
    return N
