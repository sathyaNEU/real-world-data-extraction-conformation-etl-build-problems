"""The association's own filed documents: the Watch protocol, the partner data agreements, the
partners' monthly returns, the register's earlier monthly bulletins and the data dictionary."""
import os

from reportlab import rl_config
rl_config.invariant = 1
from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import world as W

SS = getSampleStyleSheet()
BODY = ParagraphStyle("b", parent=SS["BodyText"], fontName="Times-Roman", fontSize=9.5, leading=13,
                      alignment=TA_JUSTIFY, spaceAfter=5)
H1 = ParagraphStyle("h1", parent=SS["Heading1"], fontName="Times-Bold", fontSize=13, spaceAfter=8)
H2 = ParagraphStyle("h2", parent=SS["Heading2"], fontName="Times-Bold", fontSize=10.5, spaceAfter=4,
                    spaceBefore=8)
SMALL = ParagraphStyle("s", parent=BODY, fontSize=8, leading=10.5)
CELL = ParagraphStyle("c", fontName="Times-Roman", fontSize=8, leading=10)
CELLB = ParagraphStyle("cb", fontName="Times-Bold", fontSize=8, leading=10)
TBL = TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.Color(0.45, 0.45, 0.45)),
                  ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.90, 0.90, 0.88)),
                  ("VALIGN", (0, 0), (-1, -1), "TOP"),
                  ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                  ("TOPPADDING", (0, 0), (-1, -1), 1.6), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6)])


def wrap(rows):
    return [[Paragraph(str(c), CELLB if i == 0 else CELL) for c in row] for i, row in enumerate(rows)]


def _doc(path, title):
    return SimpleDocTemplate(path, pagesize=A4, title=title, author="", subject="", creator="",
                             leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm,
                             bottomMargin=18 * mm)


def protocol(target):
    path = os.path.join(target, "protocolo_vigilancia_v4.pdf")
    d = _doc(path, "Protocolo da Vigilancia Nacional")
    codes = [["Codigo", "Acao", "Acao de confronto"]]
    for c, name, conf in W.ACTION_CODES:
        codes.append([c, name, "sim" if conf else "nao"])
    themes = [["Codigo", "Tema de disputa"]]
    for c, name, _, _, _ in W.THEMES:
        themes.append([c, name])
    f = [Paragraph("%s: protocolo da %s" % (W.ASSOCIATION, W.REGISTER), H1),
         Paragraph("Versao 4, aprovada pelo conselho de curadores em 24 de abril de 2026. As partes "
                   "1 e 5 e os anexos nao vao neste extrato.", SMALL),
         Paragraph("Parte 2 Acoes", H2),
         Paragraph("2.1 Uma acao e um tema de disputa numa localidade num dia de calendario. Varios "
                   "relatos da mesma acao, da mesma via de acesso ou de vias diferentes, sao uma acao.",
                   BODY),
         Paragraph("2.2 Os codigos de acao sao os da tabela seguinte. Sao acoes de confronto, e so "
                   "elas, as acoes dos codigos CF, RF e DP.", BODY),
         Table(wrap(codes), colWidths=[18 * mm, 95 * mm, 32 * mm], style=TBL, hAlign="LEFT"),
         Spacer(1, 5),
         Paragraph("2.3 A localidade de uma acao e a localidade do gazeteiro nacional, que pode vir "
                   "no relato com o nome atual ou com o nome anterior, e o municipio de uma acao e o "
                   "municipio que o gazeteiro atribui a essa localidade.", BODY),
         Paragraph("2.4 A data de uma acao e o dia de calendario local em que a acao comecou, tal "
                   "como o registo o fixa, e nao a data de nenhum relato.", BODY),
         Paragraph("2.5 Um relato cuja localidade o gazeteiro nao fixa nao produz acao.", BODY),
         Paragraph("2.6 O codigo que o registo guarda para uma acao e o do relato do monitor de "
                   "imprensa quando o monitor tem relato dessa acao, e o do relato da organizacao "
                   "parceira ou da linha de participacao publica nas restantes acoes.", BODY),
         Paragraph("2.7 Os temas de disputa do registo sao os da tabela seguinte. Um tema entra no "
                   "registo no dia em que a primeira acao desse tema e registada.", BODY),
         Table(wrap(themes), colWidths=[18 * mm, 127 * mm], style=TBL, hAlign="LEFT"),
         Spacer(1, 5),
         Paragraph("Parte 3 Periodo e indicadores", H2),
         Paragraph("3.2 O periodo de relato e de 1 de maio a 30 de outubro de 2026, inclusive. O "
                   "registo abriu em 1 de maio de 2026 e nao traz nenhuma acao anterior a essa data.",
                   BODY),
         Paragraph("3.3 Indicador 2, acoes de confronto: o numero das acoes registadas no periodo de "
                   "relato que foram acoes de confronto na acepcao de 2.2, ao mais proximo de "
                   "cinquenta acoes. O codigo que o classificador do monitor de imprensa atribui e "
                   "a leitura do classificador, que a mesa de codificacao mede contra a codificacao "
                   "manual. O numero que o boletim fixa e o das acoes que foram de confronto e nao "
                   "o dos codigos atribuidos.", BODY),
         Paragraph("3.4 As quinzenas de relato tem catorze dias e contam-se a partir de 1 de maio. "
                   "Os dias que sobram no fim do periodo entram na ultima quinzena, que por isso "
                   "tem quinze dias.", BODY),
         Paragraph("3.5 O boletim do periodo e levado ao conselho de curadores, que fixa com ele a "
                   "dotacao de terreno do ano seguinte.", BODY),
         Paragraph("Parte 4 Vias de acesso e codificacao", H2),
         Paragraph("4.1 As vias de acesso ao registo sao o monitor de imprensa, a rede de parceiros "
                   "e a linha de participacao publica.", BODY),
         Paragraph("4.2 O monitor de imprensa guarda todos os itens que a lista de palavras-chave "
                   "apanha, reune num unico relato os itens que respeitam a mesma acao e aceita os "
                   "relatos que o classificador fixa como acao coletiva com local e data; o "
                   "classificador atribui a cada relato aceite o seu codigo de acao, e os itens "
                   "rejeitados ficam no registo de itens com o motivo da rejeicao.", BODY),
         Paragraph("4.3 Cada parceiro entrega relatos ao abrigo do seu acordo de dados e e pago por "
                   "relato aceite, a tabela em vigor no dia em que o relato foi entregue. Os relatos "
                   "da rede de parceiros e da linha de participacao publica sao codificados por quem "
                   "os entrega e revistos pela mesa de codificacao.", BODY),
         Paragraph("4.4 O registo e extraido para o boletim de cada periodo; a extracao deste periodo "
                   "e de 6 de novembro de 2026.", BODY),
         Paragraph("4.5 A mesa de codificacao faz duas rondas de dupla codificacao por periodo de "
                   "relato. Cada ronda tem um corte, e o seu lote e tirado ao acaso dos relatos "
                   "aceites do monitor de imprensa entrados desde o corte da ronda anterior e "
                   "datados ate ao seu proprio corte; a mesa codifica-os a mao a medida que entram, "
                   "sem ver o codigo do classificador, e a ronda fecha no dia em que a mesa assina "
                   "o lote. A codificacao manual e a referencia da associacao. A ronda fica no "
                   "ficheiro de dupla codificacao, com o codigo do classificador e o da "
                   "codificacao manual onde a mesa o conseguiu fixar, o dia da codificacao manual "
                   "e a decisao da mesa nas linhas que foram a segunda codificacao.", BODY),
         Paragraph("4.6 O classificador do monitor de imprensa e o modelo que a configuracao do "
                   "monitor declara, com a data em que foi treinado e as alteracoes que teve dentro "
                   "do periodo.", BODY)]
    d.build(f)
    return path


def agreements(target):
    path = os.path.join(target, "acordos_dados_parceiros.pdf")
    d = _doc(path, "Acordos de dados com organizacoes parceiras")
    sched = [["Codigo", "Organizacao", "Regiao do acordo", "Tabela por relato aceite, EUR"]]
    for code, name, region, rate in W.PARTNERS:
        sched.append([code, name, region, "%.2f" % rate])
    f = [Paragraph("%s: acordos de dados com organizacoes parceiras" % W.ASSOCIATION, H1),
         Paragraph("Tabela em vigor desde 1 de janeiro de 2026, com a adenda de 24 de julho de 2026. "
                   "Documento do gabinete financeiro.", SMALL),
         Paragraph("Clausula 6 Pagamento", H2),
         Paragraph("6.1 Cada organizacao parceira e paga por relato aceite no registo, a tabela em "
                   "vigor no dia em que o relato foi entregue.", BODY),
         Paragraph("6.2 Os relatos entregues depois do fecho de um periodo de relato sao pagos com o "
                   "periodo em que foram entregues.", BODY),
         Paragraph("Tabela A Organizacoes e tabelas", H2),
         Table(wrap(sched), colWidths=[16 * mm, 72 * mm, 32 * mm, 36 * mm], style=TBL, hAlign="LEFT"),
         Spacer(1, 6),
         Paragraph("Adenda de 24 de julho de 2026", H2),
         Paragraph("A tabela do %s (%s) passa a %.2f EUR por relato aceite a partir de 1 de agosto "
                   "de 2026. As restantes clausulas do acordo mantem-se." %
                   (W.PARTNERS[3][1], W.AMENDED[0], W.AMENDED[2]), BODY),
         Paragraph("Assinada pelo gabinete financeiro da associacao e pela direcao da organizacao "
                   "parceira.", SMALL),
         Paragraph("Adenda de 8 de setembro de 2026", H2),
         Paragraph("Apos a fusao das suas estruturas de terreno, o %s (%s) passa a entregar sob o "
                   "codigo de entrega %s a partir de 10 de setembro de 2026. O acordo, a regiao e a "
                   "tabela mantem-se e as entregas sob os dois codigos sao do mesmo acordo." %
                   (W.PARTNERS[6][1], W.MERGED[0], W.MERGED[1]), BODY)]
    d.build(f)
    return path


def partner_returns(target, table):
    """Each partner's own monthly return: the reports it says it filed, month by month."""
    path = os.path.join(target, "relatorios_mensais_parceiros_2026.pdf")
    d = _doc(path, "Relatorios mensais das organizacoes parceiras")
    f = [Paragraph("Organizacoes parceiras: relatorios mensais de entrega, 2026" , H1),
         Paragraph("Compilado pelo gabinete financeiro a partir dos relatorios que cada organizacao "
                   "envia no fim de cada mes. Os totais sao os que cada organizacao declara.", SMALL)]
    months = ["maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro"]
    for code, name, region, _ in W.PARTNERS:
        rows = [["Mes", "Relatos entregues", "Nota da organizacao"]]
        for i, m in enumerate(months, 5):
            n, note = table.get((code, i), (0, ""))
            rows.append([m, str(n), note])
        f.append(Paragraph("%s (%s), %s" % (name, code, region), H2))
        f.append(Table(wrap(rows), colWidths=[26 * mm, 34 * mm, 94 * mm], style=TBL, hAlign="LEFT"))
        f.append(Spacer(1, 4))
    d.build(f)
    return path


def bulletins(target, series):
    """The register's earlier monthly bulletins: actions recorded and municipalities touched."""
    path = os.path.join(target, "boletins_anteriores_2026.pdf")
    d = _doc(path, "Boletins mensais da Vigilancia Nacional")
    rows = [["Mes", "Acoes registadas", "Municipios tocados", "Localidades tocadas"]]
    for m, a, mu, lo in series:
        rows.append([m, "%d" % a, "%d" % mu, "%d" % lo])
    f = [Paragraph("%s: boletins mensais da %s" % (W.ASSOCIATION, W.REGISTER), H1),
         Paragraph("Extrato dos boletins publicados de maio a setembro de 2026. O boletim mensal "
                   "conta acoes e territorio tocado; os indicadores de confronto sao do boletim de "
                   "periodo e vao pela primeira vez ao conselho de curadores de novembro.", SMALL),
         Paragraph("Quadro 1 Acoes e territorio, por mes", H2),
         Table(wrap(rows), colWidths=[30 * mm, 38 * mm, 40 * mm, 40 * mm], style=TBL, hAlign="LEFT"),
         Spacer(1, 6),
         Paragraph("Notas", H2),
         Paragraph("Uma acao e contada uma vez, qualquer que seja o numero de relatos ou de vias que "
                   "a tenham registado. Os municipios e as localidades sao os do gazeteiro nacional "
                   "na versao de 2024.", BODY),
         Paragraph("O boletim de maio cobre o mes em que o registo abriu, pelo que nao ha serie "
                   "anterior para comparacao.", BODY)]
    d.build(f)
    return path


def dictionary(target, sheets):
    """The data dictionary, one sheet per file in the extract."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    first = True
    for name, note, fields in sheets:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        ws["A1"] = name
        ws["A1"].font = Font(bold=True, size=12)
        ws["A2"] = note
        ws["A2"].font = Font(size=9, italic=True)
        for j, h in enumerate(["Campo", "Tipo", "Descricao"]):
            c = ws.cell(row=4, column=1 + j, value=h)
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E8E8E4")
        for i, (f, t, desc) in enumerate(fields, 5):
            ws.cell(row=i, column=1, value=f)
            ws.cell(row=i, column=2, value=t)
            c = ws.cell(row=i, column=3, value=desc)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 96
    path = os.path.join(target, "dicionario_dados.xlsx")
    wb.save(path)
    return path
