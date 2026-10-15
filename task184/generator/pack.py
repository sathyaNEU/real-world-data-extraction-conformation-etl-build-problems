"""Writing the pack: the two exports, the registers, the gazetteer, the beat map, the footfall files, the trading
hours, the extract notes, the dictionary, the tasking schedule and the patrol's own 2026 plan."""
import csv
import datetime as dt
import io
import os

import docs as D
import places as PL
import records as R
import sensors as S
import world as W

PLAN = "harbourgate_patrol_plan_2026.xlsx"


def write_csv(path, rows):
    with io.open(path, "w", encoding="utf-8", newline="") as fh:
        wr = csv.writer(fh, lineterminator="\r\n")
        for r in rows:
            wr.writerow(list(r))


def write_text(path, text):
    with io.open(path, "w", encoding="utf-8", newline="") as fh:
        fh.write(text.replace("\n", "\r\n"))


def plan_xlsx(path):
    """The neighbourhood team's record of the 2026 patrol, kept month by month as clause 5.1 of that season's
    agreement required: the hours worked in each commissioned slot, the officers who worked them and the slots
    stood down. Six slots a week were bought in 2026. It records what the patrol did, never what it met."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    import numpy as np
    g = np.random.default_rng([W.SEED, 9, W.NUDGE])
    wb = Workbook()
    ws = wb.active
    ws.title = "Hours worked"
    ws["A1"] = "Harbourgate hotspot patrol: hours worked by month, 2026 season"
    ws["A1"].font = Font(name="Calibri", size=12, bold=True)
    ws["A2"] = "Monthly return to the partnership under clause 5.1 of the 2026 agreement. Six slots a week."
    ws["A2"].font = Font(name="Calibri", size=10, italic=True)
    ws.append([])
    ws.append(["Month", "Day", "Slot", "Slots due", "Slots worked", "Hours worked", "Officer pairs",
               "Slots stood down", "Reason"])
    for c in ws[4]:
        c.font = Font(name="Calibri", size=11, bold=True)
        c.alignment = Alignment(horizontal="left")
    commissioned = [(3, 20), (4, 16), (4, 20), (5, 0), (5, 20), (6, 20)]
    months = ["April", "May", "June", "July", "August", "September"]
    reasons = ["aid to a public order operation", "abstraction to a search", "officer sickness",
               "abstraction to a missing person enquiry", "officer sickness", "training day",
               "aid to a football fixture", "abstraction to a search", "officer sickness"]
    for mi, month in enumerate(months):
        weeks = 4 if month in ("April", "June", "September") else 5
        for d, h in commissioned:
            due = weeks
            down = int(g.integers(0, 2)) if g.random() < 0.34 else 0
            worked = due - down
            hours = worked * W.SLOT_H - int(g.integers(0, 3))
            ws.append([month, W.DAYS[d], "%02d:00-%02d:00" % (h, (h + W.SLOT_H) % 24), due, worked,
                       hours, worked, down, reasons[(mi * 2 + d) % len(reasons)] if down else None])
    for col, wd in zip("ABCDEFGHI", (11, 11, 14, 10, 13, 13, 13, 17, 30)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A5"
    sm = wb.create_sheet("Season")
    sm.append(["Harbourgate hotspot patrol, 2026 season"])
    sm["A1"].font = Font(name="Calibri", size=12, bold=True)
    sm.append([])
    sm.append(["Slot", "Slots due", "Slots worked", "Hours worked"])
    for c in sm[3]:
        c.font = Font(name="Calibri", size=11, bold=True)
    rows = list(ws.iter_rows(min_row=5, values_only=True))
    for d, h in commissioned:
        label = "%02d:00-%02d:00" % (h, (h + W.SLOT_H) % 24)
        mine = [r for r in rows if r[1] == W.DAYS[d] and r[2] == label]
        sm.append(["%s %s" % (W.DAYS[d], label), sum(r[3] for r in mine), sum(r[4] for r in mine),
                   sum(r[5] for r in mine)])
    sm.append(["Season total", sum(r[3] for r in rows), sum(r[4] for r in rows), sum(r[5] for r in rows)])
    sm.cell(sm.max_row, 1).font = Font(name="Calibri", size=11, bold=True)
    for col, wd in zip("ABCD", (26, 11, 13, 13)):
        sm.column_dimensions[col].width = wd
    wb.properties.creator = W.PEOPLE["inspector"][0]
    wb.properties.lastModifiedBy = W.PEOPLE["inspector"][0]
    wb.properties.created = dt.datetime(2026, 10, 6, 11, 4, 0)
    wb.properties.modified = dt.datetime(2026, 10, 9, 15, 22, 0)
    wb.properties.title = "Harbourgate patrol plan 2026"
    wb.save(path)


def build(out, rows, fr, counts):
    """Write every file of the pack into `out` and return the file names by role."""
    os.makedirs(out, exist_ok=True)
    f = {}
    f["rms"] = R.RMS
    write_csv(os.path.join(out, R.RMS), R.rms_rows(rows, fr))
    f["legacy"] = R.LEGACY
    write_csv(os.path.join(out, R.LEGACY), R.legacy_rows(rows, fr))
    f["notes"] = D.NOTES
    write_text(os.path.join(out, D.NOTES), D.notes_md())
    f["dictionary"] = D.DICTIONARY
    write_csv(os.path.join(out, D.DICTIONARY), D.DICT_ROWS)
    f["rules"] = D.RULES
    D.rules_pdf(os.path.join(out, D.RULES))
    f["contract"] = D.CONTRACT
    D.contract_pdf(os.path.join(out, D.CONTRACT))
    f["schedule"] = D.SCHEDULE
    write_csv(os.path.join(out, D.SCHEDULE), D.schedule_rows())
    f["register"] = PL.REGISTER
    write_csv(os.path.join(out, PL.REGISTER), PL.register_rows())
    f["gazetteer"] = PL.GAZETTEER
    write_csv(os.path.join(out, PL.GAZETTEER), PL.gazetteer_rows())
    f["beats"] = PL.BEATS
    write_csv(os.path.join(out, PL.BEATS), PL.beat_rows())
    f["premises"] = PL.PREMISES
    write_csv(os.path.join(out, PL.PREMISES), PL.premises_rows())
    f["counts"] = S.COUNTS
    write_csv(os.path.join(out, S.COUNTS), S.count_rows(counts))
    f["availability"] = S.AVAIL
    write_csv(os.path.join(out, S.AVAIL), S.availability_rows())
    f["sensors"] = S.SENSREG
    write_csv(os.path.join(out, S.SENSREG), S.register_rows())
    f["trading"] = S.TRADING
    write_csv(os.path.join(out, S.TRADING), S.trading_rows())
    f["plan"] = PLAN
    plan_xlsx(os.path.join(out, PLAN))
    return f
