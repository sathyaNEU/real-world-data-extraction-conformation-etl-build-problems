"""Write the data files: the 2026 lane counts and queue-time sensor, the summer roster plans for 2026 and 2027, the
summer 2027 planning sheet, the mobile unit's 2026 postings, the lane register, the operators' 2026 traffic returns,
the hotel invoices and the expense claims export for the unit's cost centre, and the two files the analysis never
needs (the lane fault log and passenger feedback)."""
import csv
import datetime as dt

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font

import models as M
import profiles as PR
import sim as S
import world as W

LANE_COUNTS = "checkpoint_lane_counts_20260615_20260831.csv"
QUEUE_TIME = "queue_time_sensor_20260615_20260831.csv"
ROSTER_26 = "roster_plan_summer2026.csv"
ROSTER_27 = "roster_pattern_summer2027.xlsx"
PLANNING = "summer2027_planning_sheet.xlsx"
POSTINGS = "mobile_unit_postings_2026.csv"
REGISTER = "checkpoint_lane_register.csv"
RETURNS = "operator_traffic_returns_2026.csv"
HOTELS = "hotel_invoices_CC4109_summer2026.xlsx"
CLAIMS = "expense_claims_CC4109_2026.csv"
FAULTS = "lane_fault_log_2026.csv"
FEEDBACK = "passenger_feedback_summer2026.csv"

UNIT_CC = "CC-4109"
OPERATOR = {"CUL": "Cullinvale Airport DAC", "ARD": "Ardcarrig Airport Company", "FAR": "Faranloe Airport Ltd",
            "HAL": "Halloran Bay Airport Ltd", "ESK": "Eskerford Regional Airport", "BRN": "Brennagh Airport Ltd",
            "DRM": "Dromkeel Airport Company", "GLK": "Glenkessan Airport DAC"}
SORTED = sorted(W.CODES)


def _rng(*k):
    return np.random.default_rng([W.SEED, 11] + list(k))


# ---------------------------------------------------------------- 2026 lane counts and the queue-time sensor
def write_lane_counts(path, world):
    rows = []
    for c in SORTED:
        rows.extend(world["airports"][c]["lanes"])
    rows.sort(key=lambda r: (r[0], r[2], r[3], r[1]))
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["airport", "lane", "business_date", "interval_start", "pax_screened"])
        for c, lane, d, i, n in rows:
            w.writerow([c, lane, d.isoformat(), W.clock(i), n])
    return len(rows)


def write_queue_time(path, world):
    rows = sorted(world["waits"], key=lambda r: (r[0], r[1], r[2]))
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["airport", "business_date", "interval_start", "queue_time_min"])
        for c, d, i, v in rows:
            w.writerow([c, d.isoformat(), W.clock(i), "%.1f" % v])
    return len(rows)


# ---------------------------------------------------------------- roster plans
def write_roster(path, pattern, valid_from, valid_to):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["airport", "roster_pattern", "valid_from", "valid_to", "interval_start", "lanes_planned"])
        for c in SORTED:
            pl = PR.plan(c)
            for i in range(W.T):
                w.writerow([c, pattern + "-" + c, valid_from, valid_to, W.clock(i), int(pl[i])])
    return len(SORTED) * W.T


def write_roster_27(path):
    """The summer 2027 pattern as the airports' new rostering tool exports it: one row per 15 minutes, one column
    per checkpoint."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S27 lanes"
    ws["A1"] = "Summer 2027 roster pattern: lanes planned for the checkpoint's own officers"
    ws["A1"].font = Font(bold=True)
    ws["A2"] = "Pattern S27, valid 30 May to 11 September 2027. Exported 21 April 2027."
    ws.append([])
    ws.append(["Start"] + SORTED)
    for cell in ws[4]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    plans = {c: PR.plan(c) for c in SORTED}
    for i in range(W.T):
        ws.append([W.clock(i)] + [int(plans[c][i]) for c in SORTED])
    ws.column_dimensions["A"].width = 9
    for col in "BCDEFGHI":
        ws.column_dimensions[col].width = 7
    ws.freeze_panes = "B5"
    wb.save(path)
    return W.T


# ---------------------------------------------------------------- the summer 2027 planning sheet
def write_planning(path, design):
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    first = True
    for name, d0, d1 in W.PERIODS:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "Late June" if name == "late June" else name
        ws["A1"] = "Design day, %s to %s 2027: forecast passengers arriving at the central checkpoint, by 15 minutes" % (
            d0.strftime("%-d %B"), d1.strftime("%-d %B"))
        ws["A1"].font = bold
        ws.append([])
        ws.append(["Interval start"] + SORTED)
        for cell in ws[3]:
            cell.font = bold
            cell.alignment = Alignment(horizontal="center")
        for i in range(W.T):
            ws.append([W.clock(i)] + [int(design[c][name][i]) for c in SORTED])
        ws.append(["Day"] + [int(design[c][name].sum()) for c in SORTED])
        for cell in ws[ws.max_row]:
            cell.font = bold
        ws.column_dimensions["A"].width = 15
        for col in "BCDEFGHI":
            ws.column_dimensions[col].width = 9
        ws.freeze_panes = "B4"
    notes = wb.create_sheet("Notes")
    lines = [
        "Summer 2027 planning sheet",
        "Issued %s by Capacity Planning, %s." % (W.AS_OF.strftime("%-d %B %Y"), W.AGENCY),
        "",
        "Each period's sheet gives the design day for the eight central checkpoints: forecast passengers arriving at "
        "the checkpoint in each 15 minutes of the operating day (04:00 to 22:00), built from the operators' summer "
        "2027 schedules at forecast load factors and the show-up profiles observed in summer 2026.",
        "Periods: late June is 15 to 30 June; July and August are the calendar months.",
        "Airports by operations code: " + "; ".join("%s %s" % (c, W.NAME[c]) for c in SORTED) + ".",
        "",
        "Schedule changes behind the 2027 figures:",
        "HAL: the new based operator's two aircraft fly from 29 March 2027 (first wave 06:00 to 07:45, second wave "
        "16:45 to 18:15).",
        "DRM: the based operator has cut its Dromkeel rotations for summer 2027; the remaining banks keep their times.",
        "CUL, ARD, FAR: schedules as filed for summer 2027; growth on summer 2026 from fuller aircraft.",
        "",
        "Contact: Alana Marnan, Capacity Planning Analyst.",
    ]
    for i, t in enumerate(lines, start=1):
        if t:
            notes.cell(row=i, column=1, value=t)
    notes["A1"].font = bold
    notes.column_dimensions["A"].width = 120
    wb.save(path)


# ---------------------------------------------------------------- the unit's 2026 postings
def postings():
    """One row per posting week: reference, airport, dates, officers, officer-nights and the hotel purchase order.
    The unit's officers check in on the Sunday before a week and out on the Sunday that ends it; at Ardcarrig in
    weeks ten and eleven they stayed on."""
    officers = [24, 24, 23, 24, 24, 24, 22, 24, 24, 24, 24]
    g = _rng(7)
    po = 4113
    out = []
    for k, c in enumerate(W.POSTINGS_2026):
        if k:
            po += int(g.integers(3, 21))          # other purchase orders are raised in between
        d0 = W.SUMMER26[0] + dt.timedelta(days=7 * k)
        d1 = d0 + dt.timedelta(days=6)
        n = officers[k]
        nights = n * 7
        if c == "HAL":
            nights = (n - 2) * 7 + 2 * 3          # two officers sent home after three nights
        out.append({"ref": "P26-%02d" % (k + 1), "airport": c, "from": d0, "to": d1, "officers": n,
                    "officer_nights": nights, "po": "PO-26-%04d" % po})
    return out


def write_postings(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["posting_ref", "airport", "week_start", "week_end", "officers", "officer_nights", "hotel_po",
                    "team_leader"])
        for r in rows:
            w.writerow([r["ref"], r["airport"], r["from"].isoformat(), r["to"].isoformat(), r["officers"],
                        r["officer_nights"], r["po"], W.PEOPLE["team"][0]])
    return len(rows)


# ---------------------------------------------------------------- the lane register
EQUIP_PAX = ["CT X-ray, automated tray return", "CT X-ray, automated tray return", "Conventional X-ray, manual trays"]


def write_register(path):
    g = _rng(1)
    rows = []
    for c in SORTED:
        n = W.LANES[c]
        for k in range(1, n + 3):
            if k <= n:
                eq = EQUIP_PAX[0] if k <= max(2, n - 2) else EQUIP_PAX[2]
            else:
                eq = "Conventional X-ray, manual trays"
            since = dt.date(2015 + int(g.integers(0, 9)), int(g.integers(1, 13)), 1)
            rows.append([c, S.lane_id(k), "Central search hall, lane %d" % k, eq, since.isoformat()])
        for k, lid in enumerate(S.TEAM_IDS, start=1):
            rows.append([c, lid, "Mobile unit position %d" % k, "Transportable X-ray (mobile screening unit)",
                         "2024-05-01"])
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["airport", "lane", "position", "equipment", "in_service_from"])
        w.writerows(rows)
    return len(rows)


# ---------------------------------------------------------------- the operators' monthly returns
SEASON = {1: 0.52, 2: 0.55, 3: 0.66, 4: 0.79, 5: 0.86, 9: 0.88, 10: 0.77, 11: 0.58, 12: 0.63}


def passenger_lanes_by_month(world):
    """Passengers screened on the passenger lanes and the unit's, by airport and month (staff lanes excluded)."""
    out = {}
    for c in W.CODES:
        staff = set(S.staff_ids(c))
        acc = {}
        for (_, lane, d, i, n) in world["airports"][c]["lanes"]:
            if lane in staff:
                continue
            acc[d.month] = acc.get(d.month, 0) + n
        out[c] = acc
    return out


def write_returns(path, world):
    g = _rng(2)
    pm = passenger_lanes_by_month(world)
    rows = []
    for c in SORTED:
        july = pm[c][7]
        june_daily = pm[c][6] / 16.0
        for m in range(1, 13):
            if m == 6:
                dep = pm[c][6] + 14 * june_daily * (0.965 + 0.01 * g.random())
            elif m in (7, 8):
                dep = pm[c][m]
            else:
                days = (dt.date(2026 + (m == 12), m % 12 + 1, 1) - dt.date(2026, m, 1)).days
                dep = july / 31.0 * days * SEASON[m] * (1 + 0.015 * g.standard_normal())
            dep = int(round(dep * (1.0025 + 0.0018 * g.standard_normal())))
            arr = int(round(dep * (1.004 + 0.012 * g.standard_normal())))
            rows.append([c, OPERATOR[c], "2026-%02d" % m, dep, arr])
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["airport", "operator", "month", "departing_passengers", "arriving_passengers"])
        w.writerows(rows)
    return len(rows)


# ---------------------------------------------------------------- hotel invoices for the unit's cost centre
HOTEL = {"CUL": "Cullinvale Quayside Hotel", "ARD": "Ardcarrig Airport Lodge", "FAR": "Faranloe Bridge House",
         "HAL": "Halloran Bay Strand Hotel", "ESK": "Eskerford Mill Hotel", "BRN": "Brennagh Square Hotel",
         "DRM": "Dromkeel Court Hotel", "GLK": "Glenkessan Park Hotel"}
# per officer-night: room, breakfast (None where the rate includes it), local tax
RATES = {"CUL": (104.50, 14.20, 2.40), "ARD": (92.00, 12.80, 2.10), "FAR": (88.40, None, 1.90),
         "HAL": (96.80, 13.50, 2.20), "ESK": (84.60, 12.10, 1.80), "BRN": (79.90, None, 1.60),
         "DRM": (82.30, 11.40, 1.70), "GLK": (76.50, 11.90, 1.50)}
AUGUST_UPLIFT = {"CUL": 1.06, "ARD": 1.05, "BRN": 1.04, "GLK": 1.05}
# the Quayside was full for the unit's second Cullinvale week
OTHER_HOTEL = {"P26-07": ("Cullinvale Airport Inn", (96.00, 13.60, 2.35))}


def invoices(posts):
    """Invoice lines as the payables system holds them: every line carries the stay's officer-nights. Ardcarrig's
    weeks ten and eleven are one stay and one invoice; Halloran Bay's block booking was invoiced in full and two
    officers' unused nights credited."""
    g = _rng(3)
    stays = []
    for p in posts:
        if stays and stays[-1]["airport"] == p["airport"] and stays[-1]["to"] + dt.timedelta(days=1) == p["from"]:
            stays[-1]["to"] = p["to"]
            stays[-1]["nights"] += p["officer_nights"]
            stays[-1]["po"].append(p["po"])
            continue
        stays.append({"airport": p["airport"], "from": p["from"], "to": p["to"], "nights": p["officer_nights"],
                      "po": [p["po"]], "officers": p["officers"], "ref": p["ref"]})
    lines = []
    inv_no = {c: 20410 + 37 * k for k, c in enumerate(SORTED)}
    for s in stays:
        c = s["airport"]
        hotel, (room, bfast, tax) = OTHER_HOTEL.get(s["ref"], (HOTEL[c], RATES[c]))
        up = AUGUST_UPLIFT.get(c, 1.0) if s["from"].month == 8 else 1.0
        room = round(room * up, 2)
        billed = s["nights"] if c != "HAL" else s["officers"] * 7
        inv_no[c] += int(g.integers(3, 40))
        prefix = "".join(wd[0] for wd in hotel.replace(",", "").split()).upper()
        no = "%s-%d" % (prefix, inv_no[c])
        date = s["to"] + dt.timedelta(days=int(g.integers(1, 5)))
        items = [("Accommodation, room only" if bfast else "Accommodation, bed and breakfast", room)]
        if bfast:
            items.append(("Breakfast", bfast))
        items.append(("Local accommodation tax", tax))
        for k, (desc, rate) in enumerate(items, start=1):
            lines.append({"supplier": hotel, "invoice_no": no, "invoice_date": date, "po": s["po"][0],
                          "line": k, "desc": desc, "nights": billed, "amount": round(billed * rate, 2),
                          "airport": c})
        if c == "HAL":
            credit = billed - s["nights"]
            cn = "%s-CN%d" % (prefix, inv_no[c] + 2)
            cdate = date + dt.timedelta(days=9)
            for k, (desc, rate) in enumerate(items, start=1):
                kept = 0.10 if k == 1 else 0.0           # the block-booking terms keep a tenth of the room rate
                lines.append({"supplier": HOTEL[c], "invoice_no": cn, "invoice_date": cdate, "po": s["po"][0],
                              "line": k, "desc": desc + " (credit, unused nights)", "nights": -credit,
                              "amount": -round(credit * rate * (1 - kept), 2), "airport": c})
    return lines, stays


def write_hotels(path, lines):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AP lines"
    head = ["Supplier", "Invoice no", "Invoice date", "Cost centre", "PO", "Line", "Line description", "Nights",
            "Amount EUR"]
    ws.append(head)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in sorted(lines, key=lambda x: (x["invoice_date"], x["invoice_no"], x["line"])):
        ws.append([r["supplier"], r["invoice_no"], r["invoice_date"], UNIT_CC, r["po"], r["line"], r["desc"],
                   r["nights"], r["amount"]])
        ws.cell(row=ws.max_row, column=3).number_format = "dd/mm/yyyy"
        ws.cell(row=ws.max_row, column=9).number_format = "#,##0.00"
    widths = [28, 14, 12, 11, 13, 6, 46, 8, 12]
    for k, wdt in enumerate(widths):
        ws.column_dimensions["ABCDEFGHI"[k]].width = wdt
    ws.freeze_panes = "A2"
    wb.save(path)
    return len(lines)


# ---------------------------------------------------------------- expense claims for the unit's cost centre
SUBS_DAY = 38.20
KM_RATE = 0.3906
HOME_KM = {"CUL": 64, "ARD": 118, "FAR": 141, "HAL": 236, "ESK": 172, "BRN": 205, "DRM": 259, "GLK": 287}


def officers_list():
    g = _rng(4)
    nums = sorted({int(x) for x in g.integers(30100, 48900, 60)})[:28]
    return ["E%05d" % n for n in nums]


def claims(posts):
    """One claim per officer per posting: a day's subsistence for each day of the week, travel out on the Sunday
    the officer checks in and home on the Sunday the week ends, and receipted incidentals. Approved one to three
    weeks after the week ends."""
    g = _rng(5)
    staff = officers_list()
    out, seq = [], 61200
    for k, p in enumerate(posts):
        c = p["airport"]
        team = sorted(staff[j] for j in g.choice(len(staff), p["officers"], replace=False))
        prev_same = k > 0 and posts[k - 1]["airport"] == c
        next_same = k + 1 < len(posts) and posts[k + 1]["airport"] == c
        early = set(team[:2]) if c == "HAL" else set()
        drivers = set(team[j] for j in g.choice(len(team), 4, replace=False))
        for o in team:
            seq += int(g.integers(1, 4))
            ref = "EC26-%05d" % seq
            lines = []
            days = 3 if o in early else 7
            if not prev_same:
                out_day = p["from"] - dt.timedelta(days=1)
                if o in drivers:
                    lines.append((out_day, "MILEAGE", round(HOME_KM[c] * KM_RATE, 2)))
                else:
                    lines.append((out_day, "RAIL-BUS", round(9.40 + HOME_KM[c] * 0.071, 2)))
            for j in range(days):
                lines.append((p["from"] + dt.timedelta(days=j), "SUBS-DAY", SUBS_DAY))
            home = p["from"] + dt.timedelta(days=days - 1 if o in early else 6)
            if not next_same or o in early:
                if o in drivers:
                    lines.append((home, "MILEAGE", round(HOME_KM[c] * KM_RATE, 2)))
                else:
                    lines.append((home, "RAIL-BUS", round(9.40 + HOME_KM[c] * 0.071, 2)))
            for _ in range(int(g.integers(0, 3))):
                d = p["from"] + dt.timedelta(days=int(g.integers(0, days)))
                amt = round(float(g.choice([4.50, 6.80, 8.20, 11.60, 12.40, 14.90, 18.30, 21.70])), 2)
                lines.append((d, "INCID", amt))
                if g.random() < 0.18:                   # a taxi there and back the same day, same fare
                    lines.append((d, "INCID", amt))
            lines.sort(key=lambda x: (x[0], x[1]))
            approved = p["to"] + dt.timedelta(days=int(g.integers(9, 22)))
            for ln, (d, cat, amt) in enumerate(lines, start=1):
                out.append({"claim_ref": ref, "line": ln, "officer": o, "date": d, "category": cat, "amount": amt,
                            "approved": approved, "airport": c, "posting": p["ref"]})
    # two pre-season and post-season trips charged to the unit (briefing at headquarters, equipment return)
    for d, who, trip in ((dt.date(2026, 6, 11), staff[3], (("RAIL-BUS", 23.80), ("SUBS-DAY", SUBS_DAY))),
                         (dt.date(2026, 8, 31), staff[11], (("MILEAGE", 97.65), ("SUBS-DAY", SUBS_DAY)))):
        seq += 1
        for ln, (cat, amt) in enumerate(trip, start=1):
            out.append({"claim_ref": "EC26-%05d" % seq, "line": ln, "officer": who, "date": d, "category": cat,
                        "amount": amt, "approved": d + dt.timedelta(days=12), "airport": None, "posting": None})
    return out


def export_claims(lines):
    """The cost centre's claims export as finance keeps it: one monthly batch per approval month appended in turn,
    each batch numbering its rows on from the last; the August batch was run and appended twice."""
    out, rid = [], 880000
    for month in (6, 7, 8, 9, 10):
        batch = sorted((x for x in lines if x["approved"].month == month),
                       key=lambda x: (x["approved"], x["claim_ref"], x["line"]))
        runs = 2 if month == 8 else 1
        for _ in range(runs):
            for x in batch:
                rid += 1
                out.append(dict(x, row_id=rid))
    return out


def write_claims(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["export_row_id", "claim_ref", "line_no", "officer_no", "cost_centre", "claim_date", "category",
                    "amount_eur", "approved_on"])
        for r in rows:
            w.writerow([r["row_id"], r["claim_ref"], r["line"], r["officer"], UNIT_CC, r["date"].isoformat(),
                        r["category"], "%.2f" % r["amount"], r["approved"].isoformat()])
    return len(rows)


# ---------------------------------------------------------------- distractors
def write_faults(path, world):
    rows = []
    for c in SORTED:
        for d, lane, s, e, (code, text) in world["airports"][c]["faults"]:
            rows.append([c, S.lane_id(lane), d.isoformat(), W.clock(s), W.clock(e + 1) if e + 1 < W.T else "22:00",
                         code, text])
    rows.sort(key=lambda r: (r[2], r[0], r[3]))
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["airport", "lane", "fault_date", "lane_down_from", "lane_back_at", "fault_code", "description"])
        w.writerows(rows)
    return len(rows)


FEEDBACK_TEXT = [
    ("Queue for security went back past the shops.", (1, 2)), ("Waited about 40 minutes at security.", (1, 2)),
    ("Security staff were friendly, queue moved fine.", (4, 5)), ("Nobody told us how long the wait would be.", (1, 2)),
    ("Missed my boarding call because of the security queue.", (1, 1)), ("Fast track was closed, long wait.", (1, 3)),
    ("Trays kept jamming at the belt.", (2, 3)), ("Very quick through security this time.", (4, 5)),
    ("Queue was long but well managed.", (3, 4)), ("No signage for where the queue started.", (2, 3)),
    ("Half an hour in the queue with two small children.", (1, 1)),
    ("Staff kept us moving and were cheerful about it.", (4, 5)),
    ("The liquids bin was overflowing and nobody emptied it.", (2, 3)),
    ("Through in five minutes at six in the morning.", (5, 5)),
    ("Queue doubled back on itself and people were cutting in.", (1, 2)),
    ("Officer explained the laptop rule clearly, no fuss.", (4, 5)),
    ("Arrived two hours early and still nearly missed the flight.", (1, 1)),
    ("The queue barely moved for twenty minutes, then went quickly.", (2, 3)),
    ("Wheelchair assistance took us round the queue, very good.", (5, 5)),
    ("Had to repack my bag on the floor, no bench.", (2, 3)),
    ("Boarding card scanner at the entrance kept rejecting my phone.", (2, 3)),
    ("Far better than last summer.", (4, 5)), ("Still queuing outside the terminal door.", (1, 1)),
    ("Bag search was quick and the officer was polite.", (4, 5)),
    ("Twenty-five minutes, which for a bank holiday is fair enough.", (3, 4)),
    ("Air conditioning in the queue hall was not coping.", (2, 3)),
]



def write_feedback(path, world):
    g = _rng(6)
    rows = []
    ref = 52000
    for d in world["days"]:
        for c in SORTED:
            daily = next(x for x in world["airports"][c]["daily"] if x["date"] == d)
            lam = 0.4 + 9.0 * daily["over"] / max(1, daily["arrivals"])
            for _ in range(int(g.poisson(lam))):
                ref += int(g.integers(1, 5))
                hour = int(g.choice([5, 6, 6, 7, 7, 8, 12, 16, 17, 17, 18, 19]))
                text, (lo, hi) = FEEDBACK_TEXT[int(g.integers(0, len(FEEDBACK_TEXT)))]
                stars = int(g.integers(lo, hi + 1))
                rows.append(["FB-%d" % ref, c, d.isoformat(), "%02d:%02d" % (hour, int(g.integers(0, 60))),
                             "web form" if g.random() < 0.7 else "airport kiosk", stars, text])
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["feedback_ref", "airport", "visit_date", "visit_time", "channel", "rating", "comment"])
        w.writerows(rows)
    return len(rows)
