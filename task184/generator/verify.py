"""Independent verifier: reads only the shipped pack and the goldens and recomputes every graded figure by its own
route (the slot grid built minute by minute over each window rather than by overlap shares, the quarter taken from
the register's own street ranges, the two exports reconciled on the occurrence key, the footfall hours counted
against the supplier's fault log), then compares them with the goldens.

    python3 task184/generator/verify.py [target] [golden]"""
import csv
import datetime as dt
import io
import os
import re
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
RMS = "kcp_rms_offence_extract_2027-01-22.csv"
LEGACY = "kcp_cis_offence_extract_2025.csv"
REGISTER = "harbourgate_location_register.csv"
GAZETTEER = "kelverstoke_street_gazetteer.csv"
SCHEDULE = "harbourgate_tasking_schedule.csv"
CONTRACT = "harbourgate_patrol_tasking_contract_2027.pdf"
COUNTS = "harbourgate_footfall_hourly.csv"
FAULTS = "footfall_sensor_availability.csv"
TRADING = "harbourgate_trading_hours.csv"
PAPER = "harbourgate_patrol_slate_2027.pdf"
BOOK = "harbourgate_slot_weights.xlsx"
SEASONS = [(dt.date(2025, 4, 1), dt.date(2025, 9, 30)), (dt.date(2026, 4, 1), dt.date(2026, 9, 30))]

n = [0]


def ok(cond, what):
    n[0] += 1
    if not cond:
        raise AssertionError("verify: " + what)


def rows_of(path):
    with io.open(path, encoding="utf-8", newline="") as fh:
        r = list(csv.reader(fh))
    return [dict(zip(r[0], x)) for x in r[1:]]


def pdf_text(path):
    from pypdf import PdfReader
    return re.sub(r"\s+", " ", " ".join(pg.extract_text() for pg in PdfReader(path).pages))


def season_of(d):
    for i, (a, b) in enumerate(SEASONS):
        if a <= d <= b:
            return i
    return None


def ts(x):
    return dt.datetime.strptime(x, "%Y-%m-%d %H:%M")


def minute_grid(rows, slot_h=4):
    """Each offence's weight spread minute by minute over its window: a different arithmetic from the generator's
    overlap shares, and it has to give the same figures."""
    w = {(d, h): 0.0 for d in range(7) for h in range(0, 24, slot_h)}
    for r in rows:
        a, b = r["from"], r["to"]
        if b <= a:
            w[(a.weekday(), (a.hour // slot_h) * slot_h)] += 1.0
            continue
        mins = int((b - a).total_seconds() // 60)
        each = 1.0 / mins
        t = a
        for _ in range(mins):
            w[(t.weekday(), (t.hour // slot_h) * slot_h)] += each
            t += dt.timedelta(minutes=1)
    return w


def main(target=None, golden=None):
    target = target or os.path.join(TASK, "target")
    golden = golden or os.path.join(TASK, "golden")
    files = os.listdir(target)
    ok(len(files) >= 10, "ten files or more")
    ok(len({f.rsplit(".", 1)[1] for f in files}) >= 3, "three formats or more")

    # the contract's own words fix the population, the window and the unit
    ct = pdf_text(os.path.join(target, CONTRACT))
    ok("eight patrol slots in each week" in ct, "the contract buys eight slots")
    ok("occurred within the quarter" in ct, "the contract's quarter clause")
    classes = [c for c in ("Theft from a motor vehicle", "Theft of a pedal cycle", "Criminal damage",
                           "Burglary business", "Burglary residential") if c in ct]
    ok(len(classes) == 5, "Annex B names five classes")

    reg = rows_of(os.path.join(target, REGISTER))
    codes = {r["location_code"] for r in reg}
    ok(len(codes) == 5, "the register holds five locations")

    sched = rows_of(os.path.join(target, SCHEDULE))
    starts = sorted({int(r["slot_start_local"][:2]) for r in sched})
    ok(starts == [0, 4, 8, 12, 16, 20] and len(sched) == 42, "the schedule's 42 slots from midnight")

    # the main population, from the extract alone
    ext = rows_of(os.path.join(target, RMS))
    main_rows = []
    for r in ext:
        if r["offence_class"] not in classes or r["location_code"] not in codes:
            continue
        a, b = ts(r["occurred_from"]), ts(r["occurred_to"])
        si = season_of(a.date())
        if si is None:
            continue
        main_rows.append({"from": a, "to": b, "cls": r["offence_class"], "season": si})
    ok(len(main_rows) > 20000, "the main population is the quarter's contracted offences (%d)" % len(main_rows))
    ok(all(r["from"].minute == 0 and r["to"].minute == 0 for r in main_rows if r["to"] > r["from"]),
       "every window is on the hour")

    grid = {k: v / len(SEASONS) for k, v in minute_grid(main_rows).items()}
    order = sorted(grid.items(), key=lambda kv: -kv[1])
    slate = [s for s, _ in order[:8]]
    total = sum(v for _, v in order[:8])
    fifty = lambda x: int(round(x / 50.0) * 50)

    # the paper
    paper = pdf_text(os.path.join(golden, PAPER))
    for s in slate:
        label = "%s %02d:00 to %02d:00" % (DAYS[s[0]], s[1], (s[1] + 4) % 24)
        ok(label in paper, "the paper names %s" % label)
        ok("{:,}".format(fifty(grid[s])) in paper, "the paper's figure for %s" % label)
    ok("{:,}".format(fifty(total)) in paper, "the paper's covered total %d" % fifty(total))
    runner = order[8]
    ok("%s %02d:00 to %02d:00" % (DAYS[runner[0][0]], runner[0][1], (runner[0][1] + 4) % 24) in paper,
       "the paper names the runner-up")

    # the workbook's grid
    book = openpyxl.load_workbook(os.path.join(golden, BOOK))
    seen = 0
    for day, slot, fig, mark in book["Week"].iter_rows(min_row=2, values_only=True):
        s = (DAYS.index(day), int(slot[:2]))
        ok(fig == fifty(grid[s]), "the workbook's figure for %s %s (%d against %d)" % (day, slot, fig, fifty(grid[s])))
        ok((mark == "commissioned") == (s in set(slate)), "the workbook's mark on %s %s" % (day, slot))
        seen += 1
    ok(seen == 42, "the workbook carries all 42 slots")

    # ask 1, the two exports reconciled on the occurrence key
    leg = rows_of(os.path.join(target, LEGACY))
    xwalk = {"THEFT - FROM PERSON": "Theft from the person", "ASSAULT - NO INJURY": "Assault without injury",
             "PUBLIC ORDER OFFENCE": "Public order", "THEFT - SHOP PREMISES": "Shoplifting"}
    numbers = {}
    for r in reg:
        numbers[r["street_name"]] = (int(r["number_from"]), int(r["number_to"]))
    gaz = rows_of(os.path.join(target, GAZETTEER))
    quarter_streets = {r["street_name"]: r["street_ref"] for r in gaz if r["locality"] == "Harbourgate"}
    refs = {r["street_ref"] for r in gaz if r["locality"] == "Harbourgate"}
    primary = {r["street_ref"]: r["street_name"] for r in gaz if r["name_status"] == "primary"}

    all_names = sorted({r["street_name"] for r in gaz}, key=len, reverse=True)
    ref_of = {}
    for r in gaz:
        ref_of.setdefault(r["street_name"], []).append((r["street_ref"], r["locality"]))

    def street_of(address):
        """The street an address sits on, and the number if it carries one: the gazetteer's names matched longest
        first, so an alternative spelling resolves to the same street reference as the primary."""
        for nm in all_names:
            i = address.find(nm)
            if i < 0:
                continue
            m = re.search(r"(\d+)\s*$", address[:i])
            num = int(m.group(1)) if m else None
            cands = ref_of[nm]
            return nm, num, cands
        return None, None, []

    def in_quarter(address, code):
        if code:
            return code in codes
        nm, num, cands = street_of(address)
        if nm is None:
            return False
        local = [ref for ref, loc in cands if loc == "Harbourgate"]
        if not local:
            return False
        name = primary[local[0]]
        if name in numbers:
            lo, hi = numbers[name]
            if num is None:
                return lo == hi                     # a named building on a single-numbered street
            return lo <= num <= hi
        return True

    def location_of(address, code):
        if code:
            return [x["location_name"] for x in reg if x["location_code"] == code][0]
        nm, _num, cands = street_of(address)
        name = primary[[ref for ref, loc in cands if loc == "Harbourgate"][0]]
        return [x["location_name"] for x in reg if x["street_name"] == name][0]

    counted = {}
    keys = set()
    for r in ext:
        if not in_quarter(r["address_text"], r["location_code"]):
            continue
        si = season_of(ts(r["occurred_from"]).date())
        if si is None:
            continue
        keys.add((r["offence_class"], r["occurred_from"], r["occurred_to"], r["address_text"]))
        counted[(r["offence_class"], si)] = counted.get((r["offence_class"], si), 0) + 1
    for r in leg:
        cls = xwalk[r["offence_type"]]
        d = dt.datetime.strptime(r["occ_from_date"] + " " + r["occ_from_time"], "%d/%m/%Y %H:%M")
        si = season_of(d.date())
        if si is None or not in_quarter(r["address_text"], ""):
            continue
        frm = d.strftime("%Y-%m-%d %H:%M")
        to = (dt.datetime.strptime(r["occ_to_date"] + " " + r["occ_to_time"], "%d/%m/%Y %H:%M").strftime(
            "%Y-%m-%d %H:%M") if r["occ_to_date"] else frm)
        if (cls, frm, to, r["address_text"]) in keys:
            continue                      # the same offence, written to both systems in the parallel fortnight
        counted[(cls, si)] = counted.get((cls, si), 0) + 1
    for cls, yr, fig in book["Offences by class"].iter_rows(min_row=2, values_only=True):
        got = counted.get((cls, yr - 2025), 0)
        ok(fig == int(round(got, -1)), "ask 1 %s %d: the workbook says %d, the pack gives %d" % (cls, yr, fig, got))

    # ask 1 again, by the other reconciliation route a solver can take: the crime information system is the
    # system of record until the class migrates, the records management system from the migration date on. It has
    # to give the same count as the matched union above, offence for offence.
    migration = {}
    for line in io.open(os.path.join(target, "kcp_rms_extract_notes.md"), encoding="utf-8"):
        parts = [x.strip() for x in line.strip().strip("|").split("|")]
        if len(parts) == 2 and parts[0] in (
                "Theft from a motor vehicle", "Theft of a pedal cycle", "Criminal damage", "Burglary business",
                "Burglary residential", "Theft from the person", "Assault without injury", "Public order",
                "Shoplifting"):
            try:
                migration[parts[0]] = dt.datetime.strptime(parts[1], "%d %B %Y").date()
            except ValueError:
                pass
    ok(len(migration) == 9, "the extract notes give a migration date for every class")
    by_migration = {}
    for r in ext:
        if not in_quarter(r["address_text"], r["location_code"]):
            continue
        si = season_of(ts(r["occurred_from"]).date())
        if si is None:
            continue
        if ts(r["reported_at"]).date() < migration[r["offence_class"]]:
            continue                      # the crime information system was still the system of record
        by_migration[(r["offence_class"], si)] = by_migration.get((r["offence_class"], si), 0) + 1
    for r in leg:
        cls = xwalk[r["offence_type"]]
        d = dt.datetime.strptime(r["occ_from_date"] + " " + r["occ_from_time"], "%d/%m/%Y %H:%M")
        si = season_of(d.date())
        if si is None or not in_quarter(r["address_text"], ""):
            continue
        by_migration[(cls, si)] = by_migration.get((cls, si), 0) + 1
    for k, v in counted.items():
        ok(by_migration.get(k, 0) == v,
           "ask 1 %s %d: the migration-date route gives %d against the matched union's %d"
           % (k[0], 2025 + k[1], by_migration.get(k, 0), v))

    # ask 2, the location shares
    at, tot = {}, {}
    for r in ext:
        code = r["location_code"]
        if not in_quarter(r["address_text"], code):
            continue
        si = season_of(ts(r["occurred_from"]).date())
        if si is None:
            continue
        name = location_of(r["address_text"], code)
        at[(name, si)] = at.get((name, si), 0) + 1
        tot[si] = tot.get(si, 0) + 1
    for r in leg:
        if not in_quarter(r["address_text"], ""):
            continue
        d = dt.datetime.strptime(r["occ_from_date"] + " " + r["occ_from_time"], "%d/%m/%Y %H:%M")
        si = season_of(d.date())
        cls = xwalk[r["offence_type"]]
        frm = d.strftime("%Y-%m-%d %H:%M")
        to = (dt.datetime.strptime(r["occ_to_date"] + " " + r["occ_to_time"], "%d/%m/%Y %H:%M").strftime(
            "%Y-%m-%d %H:%M") if r["occ_to_date"] else frm)
        if si is None or (cls, frm, to, r["address_text"]) in keys:
            continue
        name = location_of(r["address_text"], "")
        at[(name, si)] = at.get((name, si), 0) + 1
        tot[si] = tot.get(si, 0) + 1
    for name_, yr, fig in book["Location shares"].iter_rows(min_row=2, values_only=True):
        si = yr - 2025
        got = 100.0 * at[(name_, si)] / tot[si]
        ok(abs(fig - round(got, 1)) < 1e-9,
           "ask 2 %s %d: the workbook says %.1f, the pack gives %.3f" % (name_, yr, fig, got))

    # ask 3, the footfall per trading hour
    trading = {}
    for r in rows_of(os.path.join(target, TRADING)):
        d = DAYS.index(r["day_of_week"])
        o = int(r["opens_local"][:2])
        c = int(r["closes_local"][:2])
        if r["closes_on"] != r["day_of_week"]:
            c += 24                       # the schedule says the day closes on the following day
        trading[d] = (o, c)

    def open_at(day, hour):
        o, c = trading[day]
        if o <= hour < min(c, 24):
            return True
        po, pc = trading[(day - 1) % 7]
        return pc > 24 and hour < pc - 24

    faults = {}
    for r in rows_of(os.path.join(target, FAULTS)):
        a = dt.datetime.fromisoformat(r["fault_from"]).date()
        b = dt.datetime.fromisoformat(r["fault_to"]).date()
        d = a
        while d < b:
            for h in range(24):
                faults.setdefault(r["sensor_id"], set()).add((d, h))
            d += dt.timedelta(days=1)
    num, den = {}, {}
    for r in rows_of(os.path.join(target, COUNTS)):
        local = dt.datetime.strptime(r["count_hour_utc"], "%Y-%m-%dT%H:%M:%SZ") + dt.timedelta(hours=1)
        si = season_of(local.date())
        if si is None or not open_at(local.weekday(), local.hour):
            continue
        if (local.date(), local.hour) in faults.get(r["sensor_id"], set()):
            continue
        num[(r["sensor_id"], si)] = num.get((r["sensor_id"], si), 0) + int(r["pedestrians"])
        den[(r["sensor_id"], si)] = den.get((r["sensor_id"], si), 0) + 1
    for sid, yr, fig in book["Footfall"].iter_rows(min_row=2, values_only=True):
        si = yr - 2025
        got = num[(sid, si)] / den[(sid, si)]
        ok(fig == int(round(got, -1)),
           "ask 3 %s %d: the workbook says %d, the pack gives %.2f" % (sid, yr, fig, got))

    print("verify: %d checks passed from the pack alone (slate %s and five more at %d a season, runner-up %s at %d; "
          "%d quarter offences of the contracted classes; the grid built minute by minute agrees with the goldens to "
          "the stated rounding)"
          % (n[0], ", ".join("%s %02d:00" % (DAYS[s[0]][:3], s[1]) for s in slate[:3]), fifty(total),
             "%s %02d:00" % (DAYS[runner[0][0]][:3], runner[0][1]), fifty(runner[1]), len(main_rows)))


if __name__ == "__main__":
    main(*(sys.argv[1:3]))
