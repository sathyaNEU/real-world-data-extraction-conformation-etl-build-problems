"""The goldens: the bridge workbook, the one-page audit plan paper and the chart."""
import datetime as dt
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt                                                   # noqa: E402
from openpyxl import Workbook                                                     # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side            # noqa: E402
from openpyxl.utils import get_column_letter                                      # noqa: E402
from reportlab import rl_config                                                   # noqa: E402

rl_config.invariant = 1
from reportlab.lib import colors                                                  # noqa: E402
from reportlab.lib.pagesizes import A4                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                    # noqa: E402
from reportlab.lib.units import mm                                                 # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import world as W                                                                  # noqa: E402

BRIDGE = "exception_population_bridge.xlsx"
PAPER = "audit_plan_paper.pdf"
CHART = "recovered_exceptions_by_unit.png"

INK = colors.HexColor("#17212b")
HEAD = PatternFill("solid", fgColor="EEF1F4")
THIN = Side(style="thin", color="C3CBD3")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
B = Font(bold=True)
R = Alignment(horizontal="right")


def _sheet(ws, header, rows, widths, money=()):
    ws.append(header)
    for c in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = B
        cell.fill = HEAD
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(header)):
        for cell in row:
            cell.border = BOX
            if cell.column in money:
                cell.number_format = "#,##0.00"
                cell.alignment = R
            elif isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"
                cell.alignment = R
    ws.freeze_panes = "A2"


def write_bridge(path, g, a1, a2, a3):
    wb = Workbook()
    ws = wb.active
    ws.title = "Bridge"
    rows = []
    for code, _o, name, _k, _c in W.UNITS:
        u = g["units"][code]
        rows.append([code, name, u["entries"], u["postings"], u["flagged"], u["graded"],
                     u["graded"] - u["flagged"]])
    rows.append(["GROUP", "%s, twelve trading units" % W.GROUP, g["entries"], g["postings"],
                 g["flagged"], g["graded"], g["graded"] - g["flagged"]])
    _sheet(ws, ["Trading unit", "Name", "Journal entries", "Postings",
                "Entries raised as exceptions", "Postings that tripped the ruleset",
                "Bridge: postings not visible in the entry count"],
           rows, [14, 34, 15, 13, 17, 19, 23])
    ws["A%d" % (ws.max_row + 2)] = ("Basis: FY26 P1 to P9, every trading unit, every live rule. Unit "
                                    "figures to the nearest hundred postings, the group figure to the "
                                    "nearest thousand, each rounded from its own unrounded value.")
    ws["A%d" % ws.max_row].font = Font(italic=True, size=9)

    ws2 = wb.create_sheet("Payables reconciliation")
    rows = []
    for code, _o, name, _k, _c in W.UNITS:
        d = a1[code]
        rows.append([code, name, d["subledger"], d["control"], d["difference"], d["grni"],
                     d["misposted"]])
    d = a1["GROUP"]
    rows.append(["GROUP", "Twelve trading units", d["subledger"], d["control"], d["difference"],
                 d["grni"], d["misposted"]])
    _sheet(ws2, ["Trading unit", "Name", "Subledger accrual at FY25 close, EUR",
                 "Purchase ledger control account, EUR", "Difference, EUR",
                 "Of which goods received not invoiced, EUR",
                 "Of which accruals posted to another unit, EUR"],
           rows, [14, 30, 18, 18, 14, 20, 20], money=(3, 4, 5, 6, 7))

    ws3 = wb.create_sheet("Till takings")
    rows = []
    for code in W.BRANCHES:
        for q in (1, 2, 3, 4):
            d = a2["%s-Q%d" % (code, q)]
            rows.append([code, W.unit_name(code), "Q%d" % q, d["takings"], d["posted"], d["gap"]])
    _sheet(ws3, ["Trading unit", "Name", "FY25 quarter", "Till takings, EUR",
                 "Sales control account, EUR", "Gap, EUR"],
           rows, [14, 32, 12, 16, 18, 14], money=(4, 5, 6))

    ws4 = wb.create_sheet("Monitoring coverage")
    rows = []
    for code, _o, name, _k, _c in W.UNITS:
        for p in range(1, 13):
            rows.append([code, name, p, a3["%s-P%02d" % (code, p)]])
    _sheet(ws4, ["Trading unit", "Name", "FY25 period",
                 "Entries in scope of the ruleset not evaluated"], rows, [14, 32, 12, 24])
    wb.properties.creator = "Group Internal Audit"
    wb.properties.lastModifiedBy = "Group Internal Audit"
    wb.properties.created = dt.datetime(2026, 7, 3, 9, 0, 0)
    wb.properties.modified = dt.datetime(2026, 7, 3, 9, 0, 0)
    wb.save(path)
    return path


def write_chart(path, g):
    codes = [u[0] for u in W.UNITS]
    flagged = [g["units"][c]["flagged"] for c in codes]
    graded = [g["units"][c]["graded"] for c in codes]
    x = range(len(codes))
    fig, ax = plt.subplots(figsize=(10.4, 5.2), dpi=150)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    w = 0.4
    b1 = ax.bar([i - w / 2 for i in x], flagged, w, label="Entries raised as exceptions",
                color="#9fb0bf", edgecolor="#6d7f8e", linewidth=0.6)
    b2 = ax.bar([i + w / 2 for i in x], graded, w, label="Postings that tripped the ruleset",
                color="#2f6f8f", edgecolor="#1d4d64", linewidth=0.6)
    for bars, vals in ((b1, flagged), (b2, graded)):
        for bar, v in zip(bars, vals):
            ax.annotate("{:,}".format(v), (bar.get_x() + bar.get_width() / 2, bar.get_height()),
                        ha="center", va="bottom", fontsize=7.4, color="#17212b")
    ax.set_xticks(list(x))
    ax.set_xticklabels(codes, rotation=30, ha="right", fontsize=8.4)
    ax.set_ylabel("Count, FY26 P1 to P9")
    ax.set_title("Postings that tripped the continuous monitoring ruleset, against the entries raised",
                 fontsize=11.4, pad=11)
    ax.legend(frameon=False, fontsize=8.8, loc="upper right")
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", color="#dde3e8", linewidth=0.7)
    ax.set_axisbelow(True)
    fig.tight_layout()
    fig.savefig(path, facecolor="white", metadata={"Software": None, "Creation Time": None})
    plt.close(fig)
    return path


def write_paper(path, g, a1):
    BODY = ParagraphStyle("b", fontName="Helvetica", fontSize=9.6, leading=13.2, textColor=INK,
                          spaceAfter=5)
    H1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=13.4, leading=17, textColor=INK,
                        spaceAfter=2)
    H2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=10, leading=13.2, textColor=INK,
                        spaceBefore=8, spaceAfter=3)
    META = ParagraphStyle("m", fontName="Helvetica", fontSize=8.1, leading=10.4,
                          textColor=colors.HexColor("#5b6671"), spaceAfter=9)
    CELL = ParagraphStyle("c", fontName="Helvetica", fontSize=8.2, leading=10, textColor=INK)
    CELLB = ParagraphStyle("cb", parent=CELL, fontName="Helvetica-Bold")
    CELLR = ParagraphStyle("cr", parent=CELL, alignment=2)
    d = SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm,
                          topMargin=18 * mm, bottomMargin=16 * mm,
                          title="Audit plan paper: population exception count",
                          author="Group Internal Audit", subject="", creator="", producer="")
    name, role = W.PEOPLE["audit"]
    rows = [[Paragraph("Trading unit", CELLB), Paragraph("Entries raised", CELLB),
             Paragraph("Postings that tripped", CELLB), Paragraph("Postings", CELLB),
             Paragraph("Per 10,000 postings", CELLB)]]
    for code, _o, nm, _k, _c in W.UNITS:
        u = g["units"][code]
        rows.append([Paragraph("%s  %s" % (code, nm), CELL),
                     Paragraph("{:,}".format(u["flagged"]), CELLR),
                     Paragraph("{:,}".format(u["graded"]), CELLR),
                     Paragraph("{:,}".format(u["postings"]), CELLR),
                     Paragraph("%.1f" % (10_000.0 * u["graded"] / u["postings"]), CELLR)])
    rows.append([Paragraph("Group", CELLB), Paragraph("{:,}".format(g["flagged"]), CELLR),
                 Paragraph("{:,}".format(g["graded"]), CELLR),
                 Paragraph("{:,}".format(g["postings"]), CELLR),
                 Paragraph("%.1f" % (10_000.0 * g["graded"] / g["postings"]), CELLR)])
    t = Table(rows, colWidths=[64 * mm, 22 * mm, 26 * mm, 22 * mm, 26 * mm], repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c3cbd3")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef1f4")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("TOPPADDING", (0, 0), (-1, -1), 2.6),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 2.6)]))
    s = [Paragraph("Sizing the FY26 audit plan: the ledger's population exception count", H1),
         Paragraph("%s, Group Internal Audit. Prepared by %s, %s. %s."
                   % (W.GROUP, name, role, W.EXTRACT_DATE.strftime("%d %B %Y")), META),
         Paragraph("The figure the plan is sized on", H2),
         Paragraph("<b>%s postings in the group's general ledger tripped the continuous monitoring "
                   "ruleset in FY26 P1 to P9.</b> The platform raised %s journal entries as exceptions "
                   "over the same window. The two are not alternative counts of the same thing: the "
                   "platform writes one verdict for a journal entry, and an entry carries anything from "
                   "one posting to %s, so the entry count is a count of vouchers and the plan needs a "
                   "count of postings."
                   % ("{:,}".format(g["graded"]), "{:,}".format(g["flagged"]),
                      "{:,}".format(g["max_items"])), BODY),
         Paragraph("By trading unit", H2), t, Spacer(1, 7),
         Paragraph("Payables reconciliation and monitoring coverage, FY25", H2),
         Paragraph("The FY25 subledger accrual reconciles to the purchase ledger control account with a "
                   "difference of EUR %s at group level, which is goods received not yet invoiced; "
                   "accruals posted to another unit net to EUR %s across the group and move four units. "
                   "Unit and quarter detail, the FY25 till reconciliation and the monitoring coverage "
                   "shortfall are in %s."
                   % ("{:,.2f}".format(a1["GROUP"]["difference"]),
                      "{:,.2f}".format(a1["GROUP"]["misposted"]), BRIDGE), BODY)]
    d.build(s)
    return path
