"""The generator's assertions on the cut pack and the truth: the three rules (the analyst test on the clean truth),
the ladder, the estimator family's closure, the fork table's rows (SF1 to SF9), the asks' devices and their
separation from the call, the Layer 1 mishandlings, and the input gates.

    python3 task479/generator/checks.py [target]"""
import csv
import json
import os
import re
import sys
import zipfile

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import bundle as B    # noqa: E402
import pack as P      # noqa: E402
import world as W     # noqa: E402

TASK = os.path.dirname(HERE)
GOLDEN = ["Arvenne", "Brisholt", "Elsmoor", "Fallowmere", "Grindell"]
DEFAULT = ["Arvenne", "Brisholt", "Corrigan Bay", "Delvane"]


class Run:
    def __init__(self):
        self.n, self.bad = 0, []

    def ok(self, cond, label):
        self.n += 1
        if not cond:
            self.bad.append(label)
            print("  FAIL", label)


def truth_estimate(x, route):
    """The routes on the generator's clean truth (no export mess at all)."""
    P_, F = x["panel"], x["first"]
    ys = np.array(W.BENCH, dtype=float)
    out = {}
    for c in W.NAMES:
        r = np.array([F[(c, y)] / P_[(c, y)] for y in W.BENCH])
        if route == "ols":
            k, a = np.polyfit(ys, r, 1)
            r25 = a + k * 2025
        else:
            r25 = r[-1]
        out[c] = r25 * P_[(c, 2025)] / F[(c, 2024)] - 1
    return sorted(c for c in out if out[c] >= W.THRESHOLD), out


def run(target, gdir, x=None, meta_path=None, fig_path=None):
    R = Run()
    x = x or W.build()
    b = B.load(target)
    # -------------------------------------------------------------- the three rules: the analyst test, executed
    easy, _ = truth_estimate(x, "flat")
    gold, _ = truth_estimate(x, "ols")
    R.ok(easy == DEFAULT, "R3a the easy route on the clean truth files the default four (%s)" % easy)
    R.ok(gold == GOLDEN, "R3b the decisive step alone on the clean truth files the golden five (%s)" % gold)
    true25 = {c: x["ratio"][(c, 2025)] * x["panel"][(c, 2025)] / x["first"][(c, 2024)] - 1 for c in W.NAMES}
    R.ok(sorted(c for c in true25 if true25[c] >= W.THRESHOLD) == GOLDEN, "R3c the Department's 2025 print qualifies "
         "the golden five")
    # -------------------------------------------------------------- the ladder on the cut pack
    sets = {r: B.qualify(b, B.estimate(b, r)) for r in B.ROUTES}
    for r in B.GOLDEN_FAMILY + ("half_step", "log_linear"):
        R.ok(sets[r] == GOLDEN, "L1 %s files the golden five" % r)
    for r in ("flat_2024", "mean_ratio", "pooled_drift", "dept_persist"):
        R.ok(sets[r] != GOLDEN, "L2 %s files a wrong set (%s)" % (r, sets[r]))
    R.ok(sets["flat_2024"] == DEFAULT, "L3 the default path's set on the pack")
    grow = {r: B.growth(b, B.estimate(b, r)) for r in B.GOLDEN_FAMILY}
    for c in b["counties"]:
        vals = {int(round(100 * grow[r][c])) for r in B.GOLDEN_FAMILY}
        R.ok(len(vals) == 1, "C1 %s growth rounds alike across the estimator family (%s)" % (c, vals))
        lo = min(abs(100 * grow[r][c] - 100 * W.THRESHOLD) for r in B.GOLDEN_FAMILY)
        R.ok(lo >= 0.45, "C2 %s sits at least 0.45 points off the floor under every family member (%.2f)" % (c, lo))
    # the margins the default path's wrong qualifiers and misses carry
    fl = B.growth(b, B.estimate(b, "flat_2024"))
    for c in ("Corrigan Bay", "Delvane"):
        R.ok(fl[c] >= 0.062 and grow["ols"][c] <= 0.045, "L4 %s in on the default path, out on the golden" % c)
    for c in ("Elsmoor", "Fallowmere", "Grindell"):
        R.ok(fl[c] <= 0.045 and grow["ols"][c] >= 0.085, "L5 %s out on the default path, in on the golden" % c)
    # the data separates the straight line from the log-linear and from the flat carry
    res = B.fit_residuals(b)
    for c in ("Elsmoor", "Fallowmere", "Grindell"):
        R.ok(res[c][1] > 5 * res[c][0], "S1 %s: the log-linear line misfits the ratio (%.2f against %.2f per cent)" % (
            c, res[c][1], res[c][0]))
    bt = B.backtest(b)
    for c in ("Corrigan Bay", "Delvane", "Elsmoor", "Fallowmere", "Grindell"):
        sign = np.sign(bt[c])
        R.ok(len(set(sign)) == 1, "S2 %s: carrying the last benchmark by the panel missed every year 2019 to 2024 the "
                                  "same way" % c)
    # -------------------------------------------------------------- the fork table
    # SF1 first print against revised history: same set and growth (ols_revised is in the family above)
    R.ok(sets["ols_revised"] == GOLDEN, "SF1 revised prior years in the ratio history: same call")
    # SF2 the 2024 base is the TY2024 first print (filed); the revised 2024 does not exist in the pack
    R.ok(all((c, 2024) not in b["revised"] for c in b["counties"]), "SF2 no revised 2024 ships")
    # SF3 resubmissions: first received instead of latest keeps the call
    R.ok(_first_submission_call(target) == GOLDEN, "SF3 the earliest submission of a resubmitted month: same call")
    # SF4 calendar year: the Department's tax year is the calendar year (filed), the Monitor's months are calendar
    # months; SF5 the county of a report through 2025 equals the register's
    regc = {r["prop_id"]: r["county"] for r in b["reg"]}
    mism = [k for k, r in b["cur"].items() if r["y"] <= 2025 and regc[r["id"]] != r["county"]]
    R.ok(not mism, "SF5 every report through 2025 carries its register county (%d differ)" % len(mism))
    legc = {}
    for (i, y, m), r in b["leg"].items():
        legc.setdefault(i, set()).add(r["county"])
    R.ok(all(len(v) == 1 for v in legc.values()), "SF5b each legacy property carries one county")
    # SF6 the floor at 6.0 per cent applied to unrounded growth: no county within 0.45 points (C2 above)
    # SF7 equal split: 1,800,000 over five divides exactly
    R.ok(B.FUND % len(GOLDEN) == 0, "SF7 the fund splits into whole dollars")
    # SF8 the Monitor's 2025 is complete: every property active in December 2025 reported every month of 2025
    months25 = {}
    for (i, y, m) in b["cur"]:
        if y == 2025:
            months25.setdefault(i, set()).add(m)
    R.ok(all(len(v) == 12 for v in months25.values()), "SF8 every reporting property has twelve months of 2025")
    # SF9 legacy county names in two cases fold to one
    raw = {r["County"] for r in _csv(os.path.join(target, P.FILES["legacy"]))}
    R.ok({c.title() for c in raw} == set(W.NAMES), "SF9 legacy county names fold to the sixteen counties")
    # -------------------------------------------------------------- ask A
    A_ok = B.ask_a(b, "register")
    A_rep = B.ask_a(b, "report")
    A_raw = B.ask_a_raw(target, "report")
    moved = [c for c in W.NAMES if c not in A_rep or "%.1f" % A_rep[c][0] != "%.1f" % A_ok[c][0]]
    moved2 = [c for c in W.NAMES if c not in A_rep or "%.2f" % A_rep[c][1] != "%.2f" % A_ok[c][1]]
    R.ok(len(moved) >= 12, "A1 the report's county code moves %d counties' occupancy" % len(moved))
    R.ok(len(moved2) >= 12, "A2 the report's county code moves %d counties' average daily rate" % len(moved2))
    raw_moved = [c for c in W.NAMES if "%.2f" % A_raw[c][1] != "%.2f" % A_rep[c][1]]
    R.ok(len(raw_moved) >= 8, "A3 resubmissions left in move %d counties' rate" % len(raw_moved))
    tA = _truth_ask_a(x)
    R.ok(all("%.1f" % tA[c][0] == "%.1f" % A_ok[c][0] and "%.2f" % tA[c][1] == "%.2f" % A_ok[c][1] for c in W.NAMES),
         "A4 the golden handling reproduces the truth exactly")
    # -------------------------------------------------------------- ask B
    Bok = B.ask_b(b, "rooms")
    Bx = B.ask_b(b, "crosswalk")
    Bn = B.ask_b(b, "name")
    Bu = B.ask_b_unlinked(b)
    mv = [c for c in W.NAMES if "%.1f" % Bx[c][0] != "%.1f" % Bok[c][0]]
    mv2 = [c for c in W.NAMES if "%.1f" % Bx[c][1] != "%.1f" % Bok[c][1]]
    R.ok(len(mv) == 13 and len(mv2) >= 11, "B1 the crosswalk alone moves %d shares and %d changes" % (len(mv), len(mv2)))
    R.ok(sum("%.1f" % Bn[c][0] != "%.1f" % Bok[c][0] for c in W.NAMES) == 13, "B2 matching by name finds no gap")
    R.ok(sum("%.1f" % Bu[c][0] != "%.1f" % Bok[c][0] for c in W.NAMES) >= 14, "B3 the unlinked read misses")
    m = B.link(b, "rooms")
    R.ok(len(m) - len(b["xwalk"]) == 13, "B4 the rooms match links all thirteen gaps uniquely")
    truth_link = {p["old_id"]: p["new_id"] for p in x["props"] if p["close"] is None or p["close"] > (2022, 3)}
    R.ok(all(truth_link[k] == v for k, v in m.items()), "B5 every link is the truth's")
    # -------------------------------------------------------------- separation: no ask device moves the call
    R.ok(B.qualify(b, B.estimate(b, "ols")) == GOLDEN and not mism, "X1 the call never reads a 2026 row or the "
         "crosswalk, and no device row falls in 2018 to 2025")
    R.ok(all(r["y"] >= 2026 for k, r in b["cur"].items() if regc[r["id"]] != r["county"]), "X2 county-code device "
         "rows are 2026 rows only")
    # -------------------------------------------------------------- Layer 1 and gates
    files = sorted(os.listdir(target))
    R.ok(len(files) >= 10, "G1 ten or more files (%d)" % len(files))
    R.ok(len({f.rsplit(".", 1)[1] for f in files}) >= 3, "G2 three or more formats")
    big = max(sum(1 for _ in open(os.path.join(target, f), encoding="utf-8")) - 1 for f in files if f.endswith(".csv"))
    R.ok(big >= 25000, "G3 a file of 25,000 rows or more (%d)" % big)
    if meta_path:
        meta = json.load(open(meta_path))
        R.ok(len(meta["distractor_files"]) >= 2 and all(d in files for d in meta["distractor_files"]),
             "G4 two or more distractors, shipped")
        R.ok(not any(re.search(r"distract|decoy|trap", f) for f in files), "G5 no file name names a role")
    for f in files:
        if f.endswith(".csv"):
            t = open(os.path.join(target, f), encoding="utf-8").read()
            R.ok("—" not in t and "–" not in t, "G6 no dashes in %s" % f)
    # dates never after the extract
    late = [r for r in _csv(os.path.join(target, P.FILES["current"])) if r["submitted_at_utc"] > "2026-10-02T06:40:00Z"]
    R.ok(not late, "G7 no submission after the extract")
    # the golden agrees with figures.json
    if fig_path and gdir:
        fig = json.load(open(fig_path))
        R.ok(fig["call"]["qualifying"] == GOLDEN, "F1 figures.json carries the golden five")
        import openpyxl
        wb = openpyxl.load_workbook(os.path.join(gdir, "tla_rising_destinations_2027.xlsx"))
        ws = wb["Rising Destinations"]
        got = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=9).value for r in range(5, 21)}
        R.ok(all(int(round(100 * got[c])) == fig["growth_pct"][c] for c in W.NAMES), "F2 the workbook's growth "
             "matches figures.json")
    print("checks: %d run, %d failed" % (R.n, len(R.bad)))
    return not R.bad


def _csv(path):
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def _first_submission_call(target):
    b = B.load(target)
    for src, path, idf, per in (("leg", P.FILES["legacy"], "Property ID", "Report Month"),
                                ("cur", P.FILES["current"], "prop_id", "period")):
        seen = {}
        for r in _csv(os.path.join(target, path)):
            if src == "leg":
                mm, yy = r[per].split("/")
                key = (r[idf], int(yy), int(mm))
                rev = float(r["Room Revenue"])
            else:
                yy, mm = r[per].split("-")
                key = (r[idf], int(yy), int(mm))
                rev = int(r["room_revenue_cents"]) / 100
            if key not in seen:
                seen[key] = rev
        for key, rev in seen.items():
            if key in b[src]:
                b[src][key] = dict(b[src][key], rev=rev)
    return B.qualify(b, B.estimate(b, "ols"))


def _truth_ask_a(x):
    s = {}
    for r in x["rep"]:
        if r["y"] == 2026 and r["m"] <= 8:
            a = s.setdefault(r["county"], [0, 0, 0.0])
            a[0] += r["avail"]
            a[1] += r["sold"]
            a[2] += r["rev"]
    return {c: (100.0 * v[1] / v[0], v[2] / v[1]) for c, v in s.items()}


if __name__ == "__main__":
    t = sys.argv[1] if len(sys.argv) > 1 else os.path.join(TASK, "target")
    ok = run(t, os.path.join(TASK, "golden"), meta_path=os.path.join(TASK, "metadata.json"),
             fig_path=os.path.join(HERE, "figures.json"))
    sys.exit(0 if ok else 1)
