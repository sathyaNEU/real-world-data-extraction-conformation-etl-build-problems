"""The independent verifier: reads only target/ (and the goldens to compare), on its own code path (pandas, no import
of the generator or bundle.py), and recomputes the call, the fourteen rates, the default path's five, both asks and
the conformed table; then reads the goldens back (the workbook's sheets and chart, the CSV) and the PDFs' filed
sentences through its own extractor.

    python3 task479/generator/verify.py [target] [golden]
"""
import json
import os
import re
import sys

import numpy as np
import pandas as pd

AS_OF = pd.Timestamp("2026-09-30")
W0, W1 = pd.Timestamp("2001-01-01"), pd.Timestamp("2026-01-01")


def _local_date(s):
    """2023-layout UTC stamps to the Central-time date: daylight time from the second Sunday of March to the first
    Sunday of November."""
    t = pd.to_datetime(s, utc=True).dt.tz_convert("America/Chicago")
    return t.dt.tz_localize(None).dt.normalize()


def load(target):
    j = lambda k: os.path.join(target, k)
    off = pd.read_csv(j("tsrfa_vertical_datum_offsets.csv"), dtype={"county_fips": str})
    off = dict(zip(off.county_fips, off.navd88_minus_ngvd29_ft))
    mem = pd.read_csv(j("tsrfa_member_counties.csv"), dtype=str)
    names = dict(zip(mem.county_fips, mem.county))
    comm = {c: f for f, cs in zip(mem.county_fips, mem.nfip_communities) for c in cs.split(";")}
    pol = pd.read_csv(j("nfip_dsa_policy_history_2026-10-02.csv"), dtype={"policy_number": str, "prior_policy_number": str,
                                                                         "county_fips": str, "community_number": str})
    pol["start"] = pd.to_datetime(pol.original_effective_date)
    end = np.where(pol.policy_status == "ACTIVE", pd.NaT,
                   np.where(pol.policy_status == "EXPIRED", pol.term_expiration_date, pol.cancellation_date))
    pol["end"] = pd.to_datetime(pd.Series(end)).fillna(pd.Timestamp("2026-10-02"))
    shift = pol.county_fips.map(off).where(pol.elevation_datum == "NGVD29", 0.0)
    pol["lfe88"] = pol.lowest_floor_elevation + shift
    pol["bfe88"] = pol.base_flood_elevation + shift
    pol["bldg_pol"] = (pol.condo_indicator != "U") & (pol.building_coverage > 0)
    c1 = pd.read_csv(j("nfip_dsa_claims_2001-2019.csv"), dtype={"CLAIM_NO": str, "POLICY_NO": str, "COUNTY_FIPS": str})
    c1 = pd.DataFrame({"claim": c1.CLAIM_NO, "policy": c1.POLICY_NO, "bid": c1.BLDG_ID, "fips": c1.COUNTY_FIPS,
                       "dol": pd.to_datetime(c1.DATE_OF_LOSS, format="%m/%d/%Y"),
                       "reported": pd.to_datetime(c1.DATE_REPORTED, format="%m/%d/%Y"),
                       "status": c1.CLAIM_STATUS, "lfe": c1.LOWEST_FLR_ELEV, "datum": c1.ELEV_DATUM,
                       "depth": c1.WATER_DEPTH_IN, "bdmg": c1.BLDG_DAMAGE_AMT, "bpay": c1.BLDG_PAID_AMT.astype(float)})
    c2 = pd.read_csv(j("nfip_dsa_claims_2020-2026.csv"), dtype={"claimNumber": str, "policyNumber": str, "countyCode": str})
    c2 = pd.DataFrame({"claim": c2.claimNumber, "policy": c2.policyNumber, "bid": c2.buildingId, "fips": c2.countyCode,
                       "dol": _local_date(c2.dateOfLoss), "reported": _local_date(c2.originalReportDate),
                       "status": c2.claimStatus, "lfe": c2.lowestFloorElevation, "datum": c2.verticalDatum,
                       "depth": c2.waterDepthInches, "bdmg": c2.buildingDamageAmount,
                       "bpay": c2.amountPaidOnBuildingClaim.astype(float)})
    cl = pd.concat([c1, c2], ignore_index=True)
    cl["lfe88"] = cl.lfe + cl.fips.map(off).where(cl.datum == "NGVD29", 0.0)
    return dict(off=off, names=names, comm=comm, pol=pol, cl=cl, target=target)


def union_days(g):
    s = g.sort_values("start")
    tot, cs, ce = 0, None, None
    for a, b in zip(s.start.clip(lower=W0), s.end.clip(upper=W1)):
        if b <= a:
            continue
        if ce is None or a > ce:
            if ce is not None:
                tot += (ce - cs).days
            cs, ce = a, b
        elif b > ce:
            ce = b
    if ce is not None:
        tot += (ce - cs).days
    return tot


def rates(L, drop_old_floor=False):
    pol, cl = L["pol"], L["cl"]
    bp = pol[pol.bldg_pol]
    cur = bp[(bp.start <= AS_OF) & (bp.end > AS_OF)].sort_values("start").groupby("nfip_building_id").tail(1)
    cur = cur[cur.bfe88.notna() & (cur.lfe88 >= cur.bfe88 - 1e-9)]
    pop = cur.set_index("nfip_building_id")[["county_fips", "lfe88"]]
    days = bp[bp.nfip_building_id.isin(pop.index)].groupby("nfip_building_id").apply(union_days)
    pop = pop.assign(days=days)
    c = cl[cl.bid.isin(pop.index) & (cl.dol >= W0) & (cl.dol < W1)]
    c = c[c.policy.isin(set(bp.policy_number))]
    c = c.join(pop[["lfe88"]].rename(columns={"lfe88": "now"}), on="bid")
    if drop_old_floor:
        keep = (c.bpay > 0) & ~(c.lfe88 < c.now - 0.45)
    else:
        keep = (c.lfe88 + c.depth / 12.0 - c.now) >= (1 / 12.0 - 1e-9)
    per_b = keep.groupby(c.bid).sum().reindex(pop.index, fill_value=0)
    pop = pop.assign(claims=per_b)
    agg = pop.groupby("county_fips").agg(claims=("claims", "sum"), days=("days", "sum"))
    agg["years"] = agg.days / 365.25
    agg["rate"] = 1000.0 * agg.claims / agg.years
    return agg.sort_values("rate"), pop


def losses(L):
    pol, cl = L["pol"], L["cl"].copy()
    cl = cl[cl.dol.dt.year == 2026].copy()
    p = pol.set_index("policy_number")
    cl["pstart"] = cl.policy.map(p.start)
    cl["pend"] = cl.policy.map(p.end)
    cl["prior"] = cl.policy.map(p.prior_policy_number)
    early = cl[(cl.dol < cl.pstart) | (cl.dol >= cl.pend)]
    key = dict(zip(cl.claim, cl.claim))
    idx = cl.set_index(["policy", "bid", "dol"]).claim
    idx = idx[~idx.index.duplicated()]
    for r in early.itertuples():
        k = (r.prior, r.bid, r.dol)
        if isinstance(r.prior, str) and k in idx.index:
            key[r.claim] = idx.loc[k]
    cl["key"] = cl.claim.map(key)
    cl = cl.sort_values("reported")
    last = cl.groupby("key").tail(1)          # the latest record of each claim carries its payments to date
    last = last[last.bpay > 0]
    out = last.groupby("fips").agg(claims=("claim", "count"), payments=("bpay", "sum"))
    out["payments"] = out.payments.round().astype(int)
    return out


def pipeline(L):
    x = os.path.join(L["target"], "tsrfa_hma_grants_export_2026-09-30.xlsx")
    pj = pd.read_excel(x, sheet_name="Projects")
    pr = pd.read_excel(x, sheet_name="Properties", dtype={"community_number": str})
    ins = pd.read_excel(x, sheet_name="Inspections")
    done = ins[(ins.inspection == "Final") & (ins.result == "Pass")].groupby("property_id").inspection_date.max()
    pr["final"] = pr.property_id.map(done)
    pr = pr.merge(pj[["project_number", "federal_cost_share_pct"]], on="project_number")
    pend = pr[(pr.property_status != "Withdrawn") & (pr.final.isna() | (pr.final > AS_OF))].copy()
    pend["share"] = (pend.approved_cost * pend.federal_cost_share_pct / 100.0).round().astype(int)
    pend["fips"] = pend.community_number.map(L["comm"])
    return pend.groupby("fips").agg(buildings=("property_id", "count"), share=("share", "sum"))


def pdf_text(path):
    from pypdf import PdfReader
    return re.sub(r"\s+", " ", " ".join(p.extract_text() or "" for p in PdfReader(path).pages))


def main(target, gdir):
    L = load(target)
    nm = L["names"]
    fails = []

    def ok(cid, cond, msg):
        print("%-4s %s  %s" % (cid, "ok  " if cond else "FAIL", msg))
        if not cond:
            fails.append(cid)
    agg, pop = rates(L)
    five = sorted(nm[f] for f in agg.index[:5])
    nxt = nm[agg.index[5]]
    ok("V1", five == sorted(["Wickmere", "Amberjack Bay", "Thornock", "Lisgrove", "Saltreach"]) and nxt == "Eastmarsh",
       "the golden five %s, next %s" % (five, nxt))
    e1, _ = rates(L, drop_old_floor=True)
    e1five = sorted(nm[f] for f in e1.index[:5])
    ok("V2", e1five == sorted(["Kettlemoor", "Dunholme", "Wickmere", "Amberjack Bay", "Thornock"]),
       "the default path's five %s" % e1five)
    a = losses(L)
    b = pipeline(L)
    fig = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "figures.json")))
    ok("V3", all(fig["components"]["rates"][nm[f]] == "%.1f" % (agg.loc[f, "rate"] + 1e-9) for f in agg.index),
       "figures.json's fourteen rates recompute")
    ok("V4", all(fig["a_losses"][nm[f]] == {"claims": int(a.loc[f, "claims"]), "payments": int(a.loc[f, "payments"])}
                 for f in nm), "ask A recomputes in every county")
    ok("V5", all(fig["b_pipeline"][nm[f]] == {"buildings": int(b.loc[f, "buildings"]), "federal_share": int(b.loc[f, "share"])}
                 for f in nm), "ask B recomputes in every county")
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(gdir, "tsrfa_2027_private_flood_pilot.xlsx"))
    ws = wb["Pilot counties"]
    got = {r[0]: (r[2], r[3]) for r in ws.iter_rows(min_row=5, max_row=18, values_only=True)}
    ok("V6", all(abs(got[nm[f]][0] - round(agg.loc[f, "rate"], 1)) < 1e-9 for f in agg.index)
       and sorted(k for k, v in got.items() if v[1] == "Pilot") == five
       and [k for k, v in got.items() if v[1] == "Next in line"] == [nxt] and len(ws._charts) == 1,
       "the workbook's county table, the marks and the chart read back")
    g2 = {r[0]: (r[2], r[3]) for r in wb["2026 losses"].iter_rows(min_row=4, max_row=17, values_only=True)}
    g3 = {r[0]: (r[2], r[3]) for r in wb["Mitigation pipeline"].iter_rows(min_row=4, max_row=17, values_only=True)}
    ok("V7", all(g2[nm[f]] == (int(a.loc[f, "claims"]), int(a.loc[f, "payments"])) and
                 g3[nm[f]] == (int(b.loc[f, "buildings"]), int(b.loc[f, "share"])) for f in nm),
       "the workbook's ask sheets read back")
    cs = pd.read_csv(os.path.join(gdir, "tsrfa_pilot_elevated_buildings_2026-09-30.csv"), dtype={"county_fips": str})
    yrs = (pop.days / 365.25).round(1)
    m = cs.set_index("nfip_building_id")
    ok("V8", len(cs) == len(pop) and (m.paid_building_claims_as_it_stands.reindex(pop.index) == pop.claims).all()
       and (abs(m.building_years.reindex(pop.index) - yrs) < 0.051).all(), "the conformed table: %d rows" % len(cs))
    g = pdf_text(os.path.join(target, "tsrfa_private_flood_placement_guidelines_2027.pdf"))
    ok("V9", "as they stand on 30 September 2026" in g and "five member counties with the lowest expected claim rates"
       in g and "unit owners' claims included" in g, "the guidelines' filed sentences read back from the PDF")
    co = pdf_text(os.path.join(target, "tsrfa_2025_flood_insurance_outreach_closeout.pdf"))
    ok("V10", "Building policies in force" in co and "Debbie Brady" in co, "the close-out reads back")
    print("verify: %d failed %s" % (len(fails), fails))
    return not fails


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    task = os.path.dirname(here)
    good = main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(task, "target"),
                sys.argv[2] if len(sys.argv) > 2 else os.path.join(task, "golden"))
    raise SystemExit(0 if good else 1)
