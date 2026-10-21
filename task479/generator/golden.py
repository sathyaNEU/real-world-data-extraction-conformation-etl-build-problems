"""The golden deliverable, computed from the cut pack through bundle.py: the Research desk's Rising Destinations 2027
workbook (the designation, the receipts restated onto the Department's basis, the Monitor's January to August 2026
by county and the panel's standing against its January 2018 draw), and generator/figures.json."""
import json
import os
import sys

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import bundle as B    # noqa: E402

FILE = "tla_rising_destinations_2027.xlsx"
INK = "1B2430"
GREEN = "2F5D4E"
PALE = "E8EFEC"
MUTED = "9AA5A0"


def compute(target):
    b = B.load(target)
    P = B.panel(b)
    est = B.estimate(b, "ols")
    g = B.growth(b, est)
    q = B.qualify(b, est)
    share = B.FUND // len(q)
    ys = np.array(B.BENCH, dtype=float)
    hist, line = {}, {}
    for c in b["counties"]:
        r = [b["first"][(c, y)] / P[(c, y)] for y in B.BENCH]
        k, a = np.polyfit(ys, r, 1)
        hist[c] = r
        line[c] = (a + k * 2024, a + k * 2025, k)
    return {"b": b, "P": P, "est": est, "g": g, "q": q, "share": share, "hist": hist, "line": line,
            "a": B.ask_a(b), "bb": B.ask_b(b)}


def write_xlsx(path, G):
    import openpyxl
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.marker import DataPoint
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    b = G["b"]
    wb = openpyxl.Workbook()
    hdr = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
    hfill = PatternFill("solid", fgColor=GREEN)
    body = Font(name="Calibri", size=10, color=INK)
    bold = Font(name="Calibri", size=10, color=INK, bold=True)
    thin = Border(bottom=Side(style="thin", color="B8C2BD"))
    top = Border(top=Side(style="thin", color=INK))
    qfill = PatternFill("solid", fgColor=PALE)

    def head(ws, row, cols, widths):
        for i, (h, w) in enumerate(zip(cols, widths), 1):
            c = ws.cell(row=row, column=i, value=h)
            c.font, c.fill = hdr, hfill
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left" if i == 1 else "right")
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[row].height = 30

    def title(ws, t, sub):
        ws["A1"] = t
        ws["A1"].font = Font(name="Calibri", size=13, bold=True, color=INK)
        ws["A2"] = sub
        ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="5B6470")

    # ---------------------------------------------------------------- 1 the designation
    ws = wb.active
    ws.title = "Rising Destinations"
    title(ws, "Five counties qualify for the 2027 Rising Destinations match: Arvenne, Brisholt, Elsmoor, Fallowmere "
              "and Grindell" if G["q"] == ["Arvenne", "Brisholt", "Elsmoor", "Fallowmere", "Grindell"] else
          "Counties qualifying for the 2027 Rising Destinations match",
          "Tax year 2025 taxable lodging receipts estimated on the Department of Revenue's basis, against tax year "
          "2024 as published in the TY2024 report. Research, 2 October 2026 extract.")
    cols = ["County", "Code", "TY2024 receipts, published ($)", "Monitor room revenue 2024 ($)",
            "Monitor room revenue 2025 ($)", "Department to Monitor ratio, 2024", "Ratio carried to 2025",
            "TY2025 receipts, estimated ($)", "Growth on TY2024", "Qualifies", "Match ($)"]
    head(ws, 4, cols, [15, 6, 17, 17, 17, 13, 13, 17, 10, 10, 12])
    code = {v: k for k, v in b["code"].items()}
    r = 5
    for c in b["counties"]:
        qual = c in G["q"]
        vals = [c, code[c], b["first"][(c, 2024)], round(G["P"][(c, 2024)]), round(G["P"][(c, 2025)]),
                G["hist"][c][-1], G["line"][c][1], round(G["est"][c], -5), round(G["g"][c], 2),
                "yes" if qual else "no", G["share"] if qual else 0]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=i, value=v)
            cell.font = bold if (qual and i in (1, 10, 11)) else body
            cell.border = thin
            if qual:
                cell.fill = qfill
            if i in (3, 4, 5, 8, 11):
                cell.number_format = "#,##0"
            elif i in (6, 7):
                cell.number_format = "0.000"
            elif i == 9:
                cell.number_format = "0%"
            if i >= 2:
                cell.alignment = Alignment(horizontal="right")
        r += 1
    ws.cell(row=r, column=1, value="State").font = bold
    for i, v, fm in ((3, sum(b["first"][(c, 2024)] for c in b["counties"]), "#,##0"),
                     (4, round(sum(G["P"][(c, 2024)] for c in b["counties"])), "#,##0"),
                     (5, round(sum(G["P"][(c, 2025)] for c in b["counties"])), "#,##0"),
                     (8, round(sum(G["est"][c] for c in b["counties"]), -5), "#,##0"),
                     (11, G["share"] * len(G["q"]), "#,##0")):
        cell = ws.cell(row=r, column=i, value=v)
        cell.font, cell.number_format, cell.border = bold, fm, top
    for i in range(1, 12):
        ws.cell(row=r, column=i).border = top
    ws.freeze_panes = "B5"
    notes_row = r + 2
    for k, t in enumerate([
        "The Department's TY2025 report is due in December 2026. Each county's 2025 receipts are its Monitor room "
        "revenue for 2025 times the Department-to-Monitor ratio carried one year along the straight line fitted to "
        "its first-print ratios for 2018 to 2024; the ratio has risen or fallen at a steady pace in every county since "
        "the panel was drawn, so the 2024 ratio is a year behind.",
        "Growth is shown to the nearest whole per cent; the 6 per cent floor is applied to the unrounded figure. The "
        "$1,800,000 fund is divided equally among the qualifying counties (program terms, section 4).",
    ]):
        c = ws.cell(row=notes_row + k, column=1, value=t)
        c.font = Font(name="Calibri", size=9, color="5B6470")
        ws.merge_cells(start_row=notes_row + k, start_column=1, end_row=notes_row + k, end_column=11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[notes_row + k].height = 40

    # the chart: growth by county, qualifying counties in the house green, the floor as a line
    cs = wb.create_sheet("Chart data")
    cs.append(["County", "Growth on TY2024 (%)", "6% floor"])
    order = sorted(b["counties"], key=lambda c: -G["g"][c])
    for c in order:
        cs.append([c, round(100 * G["g"][c], 1), 6.0])
    ch = BarChart()
    ch.type = "col"
    ch.style = 2
    ch.title = "TY2025 receipts growth on TY2024, estimated, against the 6% floor"
    ch.y_axis.title = "Per cent"
    ch.y_axis.numFmt = "0"
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = 12
    ch.y_axis.majorUnit = 2
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    ch.x_axis.tickLblSkip = 1
    ch.x_axis.txPr = None
    data = Reference(cs, min_col=2, min_row=1, max_row=len(order) + 1)
    cats = Reference(cs, min_col=1, min_row=2, max_row=len(order) + 1)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    s = ch.series[0]
    s.graphicalProperties = GraphicalProperties(solidFill=MUTED)
    s.graphicalProperties.line.noFill = True
    for i, c in enumerate(order):
        if c in G["q"]:
            pt = DataPoint(idx=i)
            pt.graphicalProperties = GraphicalProperties(solidFill=GREEN)
            pt.graphicalProperties.line.noFill = True
            s.dPt.append(pt)
    ch.gapWidth = 50
    floor = LineChart()
    floor.add_data(Reference(cs, min_col=3, min_row=1, max_row=len(order) + 1), titles_from_data=True)
    fl = floor.series[0]
    fl.graphicalProperties.line.solidFill = INK
    fl.graphicalProperties.line.width = 19050
    fl.graphicalProperties.line.dashStyle = "dash"
    fl.marker.symbol = "none"
    fl.smooth = False
    ch += floor
    ch.legend.position = "b"
    ch.height, ch.width = 10, 22
    ws.add_chart(ch, "A%d" % (notes_row + 3))

    # ---------------------------------------------------------------- 2 the ratio history
    hs = wb.create_sheet("Ratio history")
    title(hs, "The Department's receipts have outgrown or trailed the Monitor at a steady pace in every county",
          "First-print taxable lodging receipts over Monitor room revenue, tax years 2018 to 2024, and the "
          "straight line carried to 2025.")
    head(hs, 4, ["County"] + [str(y) for y in B.BENCH] + ["Change a year (line)", "2025 (line)"],
         [15] + [8] * 7 + [12, 10])
    for i, c in enumerate(b["counties"]):
        row = [c] + [round(v, 4) for v in G["hist"][c]] + [round(G["line"][c][2], 4), round(G["line"][c][1], 4)]
        for j, v in enumerate(row, 1):
            cell = hs.cell(row=5 + i, column=j, value=v)
            cell.font, cell.border = body, thin
            if j > 1:
                cell.number_format = "0.000"
    hs.freeze_panes = "B5"

    # ---------------------------------------------------------------- 3 the Monitor, January to August 2026
    ms = wb.create_sheet("Monitor Jan-Aug 2026")
    title(ms, "Lodging Monitor, January to August 2026, by county",
          "Occupancy is rooms sold over rooms available and average daily rate is room revenue over rooms sold, "
          "from the county's totals; each property in the county the register places it in.")
    head(ms, 4, ["County", "Occupancy (%)", "Average daily rate ($)"], [15, 13, 16])
    for i, c in enumerate(b["counties"]):
        o, adr = G["a"][c]
        for j, (v, fm) in enumerate(((c, None), (round(o, 1), "0.0"), (round(adr, 2), "#,##0.00")), 1):
            cell = ms.cell(row=5 + i, column=j, value=v)
            cell.font, cell.border = body, thin
            if fm:
                cell.number_format = fm
    ms.freeze_panes = "B5"

    # ---------------------------------------------------------------- 4 the panel's standing
    ps = wb.create_sheet("Panel standing")
    title(ps, "The Monitor panel against its January 2018 draw",
          "Properties reporting in January 2018 that still report in August 2026: their share of January 2018's "
          "panel rooms, and the change in their rooms since (rooms are room-nights available over the month's days).")
    head(ps, 4, ["County", "Share of January 2018 rooms (%)", "Change in their rooms to August 2026 (%)"],
         [15, 18, 20])
    for i, c in enumerate(b["counties"]):
        sh, chg = G["bb"][c]
        for j, (v, fm) in enumerate(((c, None), (round(sh, 1), "0.0"), (round(chg, 1), "0.0")), 1):
            cell = ps.cell(row=5 + i, column=j, value=v)
            cell.font, cell.border = body, thin
            if fm:
                cell.number_format = fm
    ps.freeze_panes = "B5"

    # ---------------------------------------------------------------- 5 notes
    ns = wb.create_sheet("Notes")
    ns.column_dimensions["A"].width = 120
    for i, t in enumerate([
        "Sources: Tarrowind Department of Revenue, Taxable Lodging Receipts by County, Tax Years 2018 to 2024 "
        "(first prints, and each report's revised prior year); Lodging Monitor reports from the legacy system "
        "(January 2018 to March 2022) and InnSight (April 2022 to September 2026), extract of 2 October 2026; the "
        "InnSight property register and the April 2022 id crosswalk.",
        "Monitor reports: where a property resubmitted a month, the later submission is used. Legacy county names are "
        "read in one case. InnSight revenue is in cents.",
        "From April 2026 the reports of the properties Harbor & Pine Hospitality took over carry its office county's "
        "code; their figures are placed in the county the register gives.",
        "Thirteen properties reporting in March 2022 are missing from the April 2022 crosswalk; each is the one "
        "InnSight property in its county that first reported in April 2022 outside the crosswalk with the same rooms, "
        "under a new name.",
        "September 2026 is not complete in the extract and is not used.",
    ]):
        c = ns.cell(row=1 + i, column=1, value=t)
        c.font = body
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ns.row_dimensions[1 + i].height = 32
    for w in wb.worksheets:
        w.sheet_view.showGridLines = False
        w.page_setup.orientation = "landscape"
        w.page_setup.fitToWidth = 1
        w.page_setup.fitToHeight = 0
        w.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "4:4"
    wb.save(path)


def figures(G):
    b = G["b"]
    return {
        "call": {"qualifying": G["q"], "match_each": G["share"], "fund": B.FUND},
        "growth_pct": {c: int(round(100 * G["g"][c])) for c in b["counties"]},
        "ask_a": {c: {"occupancy_pct": "%.1f" % G["a"][c][0], "adr_usd": "%.2f" % G["a"][c][1]}
                  for c in b["counties"]},
        "ask_b": {c: {"rooms_share_pct": "%.1f" % G["bb"][c][0], "rooms_change_pct": "%.1f" % G["bb"][c][1]}
                  for c in b["counties"]},
        "default_path": {c: int(round(100 * v)) for c, v in B.growth(b, B.estimate(b, "flat_2024")).items()},
        "conventions": {"threshold_pct": 6, "years": [2018, 2020, 2021, 2022, 2024, 2025, 2026, 2027],
                        "counts": [3, 7, 13, 16, 30], "months_2026": "January to August", "rooms_day": 31},
    }


def main(target, gdir, fig_path):
    os.makedirs(gdir, exist_ok=True)
    G = compute(target)
    write_xlsx(os.path.join(gdir, FILE), G)
    import postprocess as PP
    PP.normalize_ooxml(os.path.join(gdir, FILE), "Samantha Walter", "Samantha Walter", "2026-10-05T19:12:00Z",
                       "2026-10-06T14:27:00Z", excel=True, company="Tarrowind Lodging Association")
    with open(fig_path, "w") as fh:
        json.dump(figures(G), fh, indent=1, sort_keys=True)
        fh.write("\n")
    return G
