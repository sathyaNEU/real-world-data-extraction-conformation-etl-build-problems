"""The golden's reading of the shipped pack: the line list's cases in the window by barangay and age band,
their history, and the CBMS tables. build.py and checks.py use it; verify.py has its own reader."""
import csv
import datetime as dt
import os
import re

import numpy as np
from openpyxl import load_workbook

from common import mw_of, TARGET

LINE_LIST = "dengue_line_list_2022_2027.csv"
CBMS = "cbms_2025_dalisay_city.xlsx"


def norm(s):
    s = s.strip().lower().replace("ñ", "n")
    s = re.sub(r"^brgy\.?\s*", "", s)
    s = s.replace("(city proper)", "").replace("pob.", "poblacion")
    s = re.sub(r"\bsto\.?\s", "santo ", s)
    s = s.replace("-", " ")
    return re.sub(r"\s+", " ", s).strip()


def parse_date(s):
    s = (s or "").strip()
    if not s:
        return None
    if re.match(r"\d{4}-\d{2}-\d{2}$", s):
        return dt.date.fromisoformat(s)
    d, m, y = s.split("/")
    return dt.date(int(y), int(m), int(d))


def purok_no(addr):
    m = re.search(r"purok\s*(\d+)", addr.lower())
    return int(m.group(1)) if m else None


def read_cbms(path=None):
    path = path or os.path.join(TARGET, CBMS)
    wb = load_workbook(path, data_only=True)
    t1, t3, t4 = wb["Table 1"], wb["Table 3"], wb["Table 4"]

    def num(v):
        return 0 if v in (None, "-", "") else int(v)
    names, bands = [], []
    for row in t1.iter_rows(min_row=6, values_only=True):
        if row[0] == "Dalisay City":
            break
        if not row[0] or row[1] is None:
            continue
        names.append(row[0])
        ages = [num(v) for v in row[2:16]]
        bands.append([sum(ages[0:3]), sum(ages[3:13]), ages[13]])
    emp = {}
    for row in t3.iter_rows(min_row=6, values_only=True):
        if row[0] and row[0] != "Dalisay City" and row[1] is not None:
            emp[row[0]] = num(row[2])
    hdr = [c for c in next(t4.iter_rows(min_row=5, max_row=5, values_only=True))]
    cols = hdr[2:]
    W = {}
    for row in t4.iter_rows(min_row=6, values_only=True):
        if not row[0] or row[1] is None:
            continue
        W[row[0]] = {cols[k]: num(v) for k, v in enumerate(row[2:2 + len(cols)])}
    return names, np.array(bands, dtype=float), emp, W, cols


def read_line_list(path=None):
    path = path or os.path.join(TARGET, LINE_LIST)
    with open(path, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def city_cases(rows, canon):
    """The city's cases: latest status per case number, not-a-case and non-residents out, one case per
    patient notified twice."""
    latest = {}
    for r_ in rows:
        latest[r_["case_no"]] = r_          # the export runs in report order, so the last row is the latest
    cases = []
    for no, r_ in latest.items():
        if r_["case_classification"] in ("Not a case", "X"):
            continue
        if r_["city_municipality"].strip() != "Dalisay City":
            continue
        b = norm(r_["barangay"])
        if b not in canon:
            raise KeyError("unmapped barangay %r" % r_["barangay"])
        name = canon[b]
        p = purok_no(r_["street_purok"])
        if name == "Santo Tomas" and p is not None and p >= 5:
            name = "Santo Tomas Norte"
        onset = parse_date(r_["date_onset"]) or parse_date(r_["date_consult"])
        cases.append({"case_no": no, "barangay": name, "age": int(r_["age_years"]), "onset": onset,
                      "initials": r_["patient_initials"], "sex": r_["sex"], "dob": r_["date_of_birth"],
                      "consult": parse_date(r_["date_consult"]), "reported": parse_date(r_["date_reported"])})
    seen = set()
    out = []
    for c in sorted(cases, key=lambda c: (c["reported"], c["case_no"])):
        key = (c["initials"], c["sex"], c["age"], c["barangay"], c["onset"])
        if key in seen:
            continue
        seen.add(key)
        out.append(c)
    return out


def band(age):
    return 0 if age < 15 else (1 if age < 65 else 2)


def aggregates(cases, names, weeks=(37, 40), year=2027):
    idx = {n: k for k, n in enumerate(names)}
    res = np.zeros((len(names), 3))
    hist = np.zeros((len(names), 3))
    for c in cases:
        y, w = mw_of(c["onset"])
        i = idx[c["barangay"]]
        if y == year and weeks[0] <= w <= weeks[1]:
            res[i, band(c["age"])] += 1
        if y <= 2026:
            hist[i, band(c["age"])] += 1
    return res, hist


def load(target=None):
    t = target or TARGET
    names, pop, emp_d, W_d, cols = read_cbms(os.path.join(t, CBMS))
    canon = {norm(n): n for n in names}
    rows = read_line_list(os.path.join(t, LINE_LIST))
    cases = city_cases(rows, canon)
    res, hist = aggregates(cases, names)
    n = len(names)
    emp = np.array([emp_d[nm] for nm in names], dtype=float)
    W = np.zeros((n, n + 1))
    for i, nm in enumerate(names):
        rowd = W_d[nm]
        for j, nj in enumerate(names):
            W[i, j] = rowd.get(nj, 0)
        W[i, n] = rowd.get("Outside Dalisay City", 0)
        W[i] /= max(W[i].sum(), 1)
    return {"names": names, "pop": pop, "emp": emp, "W": W, "res": res, "hist": hist, "cases": cases,
            "rows": rows}
