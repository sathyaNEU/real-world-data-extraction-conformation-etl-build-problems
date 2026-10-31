"""Build task181's pack: world 0, the shipped files under target/, figures.json and the goldens."""
import csv
import datetime as dt
import json
import os
import shutil
import sys

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import TARGET, GOLDEN, HERE, write_csv, normalise_zip, set_mtime, mw_of, mw_start, TRIGGER  # noqa
import world as Wd  # noqa
import linelist  # noqa
import asks  # noqa
import docs  # noqa
import design  # noqa
import reader  # noqa
import routes as R  # noqa

WORLD_SEED = 3
FILES = {
    "line_list": "dengue_line_list_2022_2027.csv",
    "cbms": "cbms_2025_dalisay_city.xlsx",
    "plan": "dalisay_dengue_plan_2027_section4.pdf",
    "register": "barangays_psgc_2025.csv",
    "dmc": "dmc_dengue_inpatient_census_2027.csv",
    "ddh": "ddh_dengue_admissions_2027.xlsx",
    "larval": "larval_survey_2027_09.csv",
    "lab": "cho_lab_dengue_results_2027.csv",
    "runlog": "cho_lab_run_log_2027.json",
    "vcu": "vcu_operations_2022_2027.csv",
    "rain": "rainfall_city_gauge_2022_2027.csv",
    "bulletin": "lbilhz_dengue_update_mw39_2027.pdf",
}
DISTRACTORS = ["vcu", "rain", "bulletin"]


def world():
    cases, tables, rates = Wd.simulate_cases(WORLD_SEED)
    people = Wd.assign_attributes(cases, WORLD_SEED)
    ev = design.evaluate(people, tables)
    assert not ev["fails"], ev["fails"]
    return people, tables, ev


def write_ddh(path, header, sheets):
    wb = Workbook()
    first = True
    for (y, w) in sorted(sheets):
        if (y, w) < (2027, 30):
            continue
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "MW%02d" % w
        a = mw_start(y, w)
        b = a + dt.timedelta(days=6)
        ws["A1"] = "DALISAY DOCTORS' HOSPITAL"
        ws["A1"].font = Font(bold=True)
        ws["A2"] = "Dengue admissions log, morbidity week %d (%s to %s)" % (w, a.strftime("%d %b"), b.strftime("%d %b %Y"))
        ws.append([])
        ws.append(header)
        for row in sheets[(y, w)]:
            ws.append(row)
        ws.column_dimensions["D"].width = 30
        ws.column_dimensions["H"].width = 28
    wb.properties.creator = "DDH Medical Records"
    wb.properties.lastModifiedBy = "DDH Medical Records"
    fixed = dt.datetime(2027, 10, 11, 7, 40, 0)
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.save(path)
    normalise_zip(path)


def build():
    if os.path.isdir(TARGET):
        shutil.rmtree(TARGET)
    os.makedirs(TARGET)
    people, tables, ev = world()
    popn, grp, W, workers = tables
    header, rows, persons, first_no = linelist.build_rows(people)
    p = lambda k: os.path.join(TARGET, FILES[k])  # noqa
    write_csv(p("line_list"), header, rows)
    docs.cbms_xlsx(p("cbms"), grp, W, workers)
    docs.plan_pdf(p("plan"))
    docs.register_csv(p("register"))
    stays = asks.hospital_stays(persons)
    h, drows, dmc_truth = asks.dmc_census(stays)
    write_csv(p("dmc"), h, drows)
    oh, sheets, ddh_truth = asks.ddh_sheets(stays)
    write_ddh(p("ddh"), oh, sheets)
    lh, lrows, larval_truth = asks.larval_survey()
    write_csv(p("larval"), lh, lrows)
    labh, labrows, runlog, lab_truth = asks.laboratory(persons)
    write_csv(p("lab"), labh, labrows)
    with open(p("runlog"), "w", encoding="utf-8", newline="\n") as f:
        json.dump({"laboratory": "City Health Office Laboratory, Dalisay City", "exported": "2027-10-11T07:55:00+08:00",
                   "runs": runlog}, f, indent=1, ensure_ascii=False)
        f.write("\n")
    set_mtime(p("runlog"))
    docs.operations_log(p("vcu"))
    docs.rainfall_csv(p("rain"))
    y27 = sum(1 for q in people if q["year"] == 2027 and q["week"] <= 39)
    y26 = sum(1 for q in people if q["year"] == 2026 and q["week"] <= 39)
    docs.bulletin_pdf(p("bulletin"), int(round(y27 * 0.97)), y26)
    truths = {"dmc": dmc_truth, "ddh": ddh_truth, "larval": larval_truth, "lab": lab_truth}
    return people, tables, ev, persons, truths


# ------------------------------------------------------------------ the golden's figures from the pack
def golden_main():
    L = reader.load()
    res, hist, pop, emp, W = L["res"], L["hist"], L["pop"], L["emp"], L["W"]
    names = L["names"]
    inv = R.inversion(res, hist, pop, emp, W)
    family = {k: R.flagged(f(res, hist, pop, emp, W)) for k, f in R.ROUTES_STEP.items()}
    gold = R.flagged(inv)
    assert all(v == gold for v in family.values()), family
    default = R.flagged(R.residence(res, hist, pop, emp, W))
    return L, names, gold, default, inv, family


def _dt(s):
    return dt.datetime.strptime(s, "%d/%m/%Y %H:%M")


def golden_dmc(path):
    """DMC admissions: a stay is a run of a patient's ward segments, each starting the minute the last ended
    with no disposition; a segment after a disposition, or after a gap, starts a new admission."""
    with open(path, encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    by_no = {}
    for r_ in rows:
        by_no.setdefault(r_["hospital_no"], []).append(r_)
    counts = {w: 0 for w in asks.ASK_WEEKS_HOSP}
    for hn, segs in by_no.items():
        segs.sort(key=lambda r_: _dt(r_["ward_in"]))
        prev = None
        for r_ in segs:
            t_in = _dt(r_["ward_in"])
            continuing = (prev is not None and prev["disposition"] == "" and prev["ward_out"] != ""
                          and _dt(prev["ward_out"]) == t_in)
            if not continuing:
                y, w = mw_of(t_in)
                if y == 2027 and w in counts:
                    counts[w] += 1
            prev = r_
    return counts


def dmc_transfers_out(path):
    with open(path, encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    return [(_dt(r_["ward_out"]), int(r_["age"]), r_["sex"], r_["barangay"].upper(), r_["city_municipality"].upper())
            for r_ in rows if r_["disposition"] == "Transferred to other hospital"]


def golden_ddh(path, dmc_path):
    """DDH admissions by the date admitted, a patient moved in from DMC counted at DMC: matched by age, sex
    and address, admitted at DDH within eight hours of leaving DMC."""
    outs = dmc_transfers_out(dmc_path)
    wb = load_workbook(path, data_only=True)
    counts = {w: 0 for w in asks.ASK_WEEKS_HOSP}
    used = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=5, values_only=True):
            if not row[0]:
                continue
            t = dt.datetime.strptime(row[4] + " " + row[5], "%m/%d/%Y %I:%M %p")
            brgy, town = [x.strip().upper() for x in row[3].rsplit(",", 1)]
            match = None
            for k, (t_out, age, sex, b, tw) in enumerate(outs):
                if k in used:
                    continue
                if age == int(row[1]) and sex == row[2] and b == brgy and tw == town and \
                        dt.timedelta(0) < t - t_out <= dt.timedelta(hours=8):
                    match = k
                    break
            if match is not None:
                used.add(match)
                continue
            y, w = mw_of(t.date())
            if y == 2027 and w in counts:
                counts[w] += 1
    return counts


def golden_larval(path):
    with open(path, encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    houses = {}
    for r_ in rows:
        key = (r_["barangay"], r_["purok"], r_["household_code"])
        pos = (r_["larvae"] == "Y" or r_["pupae"] == "Y") if r_["team"] == "VCU" else int(r_["containers_positive"] or 0) > 0
        houses[key] = houses.get(key, False) or pos
    out = {}
    for b in sorted(set(k[0] for k in houses)):
        hs = [v for k, v in houses.items() if k[0] == b]
        out[b] = round(100.0 * sum(hs) / len(hs), 1)
    return out


def golden_lab(path, logpath):
    log = json.load(open(logpath, encoding="utf-8"))
    rejected = {r_["run_id"] for r_ in log["runs"] if r_["qc"] == "Rejected"}
    with open(path, encoding="utf-8") as f:
        rows = [r_ for r_ in csv.DictReader(f) if r_["run_id"] not in rejected]
    first = {}
    for r_ in sorted(rows, key=lambda r_: (dt.datetime.strptime(r_["collected_on"], "%d/%m/%Y"), r_["specimen_id"])):
        who = (r_["patient_initials"], r_["date_of_birth"], r_["sex"])
        if r_["result"] == "Positive" and who not in first:
            first[who] = r_
    counts = {(t, w): 0 for t in ("NS1", "IgM") for w in asks.ASK_WEEKS_LAB}
    for r_ in first.values():
        y, w = mw_of(dt.datetime.strptime(r_["collected_on"], "%d/%m/%Y").date())
        t = "NS1" if r_["test"].startswith("NS1") else "IgM"
        if y == 2027 and w in asks.ASK_WEEKS_LAB:
            counts[(t, w)] += 1
    return counts


def figures(truths):
    L, names, gold, default, inv, family = golden_main()
    p = lambda k: os.path.join(TARGET, FILES[k])  # noqa
    dmc = golden_dmc(p("dmc"))
    ddh = golden_ddh(p("ddh"), p("dmc"))
    larv = golden_larval(p("larval"))
    lab = golden_lab(p("lab"), p("runlog"))
    assert dmc == truths["dmc"], (dmc, truths["dmc"])
    assert ddh == truths["ddh"], (ddh, truths["ddh"])
    assert larv == truths["larval"], (larv, truths["larval"])
    assert lab == truths["lab"], (lab, truths["lab"])
    fig = {
        "call": {"hotspots": sorted(names[i] for i in gold),
                 "trigger": TRIGGER, "window": "morbidity weeks 37 to 40, 2027", "response": "weeks 41 to 44"},
        "default_path": sorted(names[i] for i in default),
        "dmc_admissions": {str(w): dmc[w] for w in asks.ASK_WEEKS_HOSP},
        "ddh_admissions": {str(w): ddh[w] for w in asks.ASK_WEEKS_HOSP},
        "house_index": larv,
        "first_positive_ns1": {str(w): lab[("NS1", w)] for w in asks.ASK_WEEKS_LAB},
        "first_positive_igm": {str(w): lab[("IgM", w)] for w in asks.ASK_WEEKS_LAB},
        "weeks": {"window": [37, 40], "response": [41, 44], "admissions": list(asks.ASK_WEEKS_HOSP),
                  "laboratory": list(asks.ASK_WEEKS_LAB)},
        "plan_sections": ["4.1", "4.2", "4.3", "4.4", "4.5", "4.6"],
        "working_age_band": [15, 64],
        "shares_at_work": [0.5, 0.6, 0.7, 0.8, 1.0],
        "over_trigger_by_home_address": len(default),
        "placement_by_inversion": {names[i]: round(float(inv[i]), 1) for i in range(len(names))},
        "residence_counts": {names[i]: int(L["res"][i].sum()) for i in range(len(names))},
    }
    with open(os.path.join(HERE, "figures.json"), "w", encoding="utf-8") as f:
        json.dump(fig, f, indent=1, ensure_ascii=False, sort_keys=True)
        f.write("\n")
    return fig, L


COVERS = {
    "line_list": "dengue notifications to the city surveillance unit, January 2022 to the extract of 11 October 2027",
    "cbms": "Community-Based Monitoring System 2025 tables for the city: population, employment, place of work",
    "plan": "the City Health Office's dengue plan for 2027, section 4 (surveillance and the hotspot response)",
    "register": "the city's barangays with their codes, districts, areas and creation notes",
    "dmc": "Dalisay Medical Center's dengue inpatient census, late July to the extract",
    "ddh": "Dalisay Doctors' Hospital's dengue admissions log by week of entry, late July to the extract",
    "larval": "the household larval survey of September 2027",
    "lab": "the City Health Office laboratory's dengue test results, July to October 2027",
    "runlog": "the City Health Office laboratory's ELISA run log with quality-control status",
    "vcu": "the Vector Control Unit's activity log, 2022 to August 2027",
    "rain": "daily rainfall at the city gauge, 2022 to October 2027",
    "bulletin": "the inter-local health zone's dengue update for weeks 1 to 39 of 2027",
}


def metadata():
    meta = {
        "task": "task181",
        "decision": "Which barangays of Dalisay City receive the dengue hotspot response for morbidity weeks 41 to 44 "
                    "of 2027, with the response figures for the hospitals, the larval survey and the laboratory.",
        "deliverables": ["dengue_hotspots_mw41_44_2027.pdf", "response_figures_2027.csv"],
        "distractor_files": ["target/" + FILES[k] for k in DISTRACTORS],
        "files": {("target/" + FILES[k]): {"source": "synthetic, generated by generator/build.py",
                                            "date": "2027-10-11", "license": "CC0-1.0", "covers": COVERS[k]}
                  for k in FILES},
    }
    path = os.path.join(os.path.dirname(HERE), "metadata.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(meta, f, indent=1, ensure_ascii=False)
        f.write("\n")


if __name__ == "__main__":
    import deliver
    people, tables, ev, persons, truths = build()
    fig, L = figures(truths)
    deliver.make(fig, L)
    metadata()
    print("hotspots", fig["call"]["hotspots"])
    print("default", fig["default_path"])
    print("rows in the line list", len(L["rows"]))
