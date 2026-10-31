"""Independent verifier: reads only the shipped bundle (target/ and the goldens), on its own code path, and
recomputes the call and every graded figure. No generator module is imported.

    python3 task181/generator/verify.py [bundle_dir]
"""
import datetime as dt
import json
import os
import re
import sys
import unicodedata

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from pypdf import PdfReader
from scipy.optimize import nnls

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
OUT = []


def ok(label, cond, detail=""):
    OUT.append((label, bool(cond), detail))


MW1 = {2021: dt.date(2021, 1, 3), 2022: dt.date(2022, 1, 2), 2023: dt.date(2023, 1, 1), 2024: dt.date(2023, 12, 31),
       2025: dt.date(2024, 12, 29), 2026: dt.date(2026, 1, 4), 2027: dt.date(2027, 1, 3), 2028: dt.date(2028, 1, 2)}


def week(d):
    for y in sorted(MW1, reverse=True):
        if d >= MW1[y]:
            return y, (d - MW1[y]).days // 7 + 1
    raise ValueError(d)


def canon(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower().strip()
    s = re.sub(r"^brgy\.?\s*", "", s)
    s = s.replace("(city proper)", "").replace("pob.", "poblacion").replace("-", " ")
    s = re.sub(r"\bsto\.?\s", "santo ", s)
    return " ".join(s.split())


def todate(s):
    s = str(s).strip()
    if not s or s == "nan":
        return None
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
        return dt.date(int(s[:4]), int(s[5:7]), int(s[8:]))
    d, m, y = s.split("/")
    return dt.date(int(y), int(m), int(d))


def cbms(path):
    wb = load_workbook(path, data_only=True)
    t1 = [r for r in wb["Table 1"].iter_rows(values_only=True)]
    names, pop = [], []
    for r in t1[5:]:
        if r[0] == "Dalisay City":
            break
        if r[0] and isinstance(r[1], int):
            v = [0 if x in (None, "-") else int(x) for x in r[2:16]]
            names.append(r[0])
            pop.append([sum(v[:3]), sum(v[3:13]), v[13]])
    t3 = {r[0]: r[2] for r in wb["Table 3"].iter_rows(values_only=True) if r[0] in names}
    rows4 = list(wb["Table 4"].iter_rows(values_only=True))
    cols = list(rows4[4][2:])
    W = np.zeros((len(names), len(names) + 1))
    for r in rows4[5:]:
        if r[0] in names and isinstance(r[1], int):
            i = names.index(r[0])
            for k, c in enumerate(cols):
                v = 0 if r[2 + k] in (None, "-") else int(r[2 + k])
                j = names.index(c) if c in names else len(names)
                W[i, j] += v
    W = W / W.sum(1, keepdims=True)
    return names, np.array(pop, float), np.array([t3[n] for n in names], float), W


def main(bundle):
    T = os.path.join(bundle, "target") if os.path.isdir(os.path.join(bundle, "target")) else bundle
    fig = json.load(open(os.path.join(HERE, "figures.json"), encoding="utf-8"))
    names, pop, emp, W = cbms(os.path.join(T, "cbms_2025_dalisay_city.xlsx"))
    ok("xlsx read: CBMS Table 1 lists the barangays (positive control: San Roque)", "San Roque" in names)
    key = {canon(n): n for n in names}
    ll = pd.read_csv(os.path.join(T, "dengue_line_list_2022_2027.csv"), dtype=str, keep_default_na=False)
    ok("csv read: the line list has 25,000 rows or more", len(ll) >= 25000, str(len(ll)))
    last = ll.groupby("case_no", sort=False).tail(1)
    last = last[~last["case_classification"].isin(["Not a case", "X"])]
    last = last[last["city_municipality"].str.strip() == "Dalisay City"].copy()
    last["b"] = [key[canon(b)] for b in last["barangay"]]
    pk = last["street_purok"].str.extract(r"(?i)purok\s*(\d+)")[0].fillna("0").astype(int)
    last.loc[(last["b"] == "Santo Tomas") & (pk >= 5), "b"] = "Santo Tomas Norte"
    last["d"] = [todate(o) or todate(c) for o, c in zip(last["date_onset"], last["date_consult"])]
    last["rep"] = [todate(x) for x in last["date_reported"]]
    last["age"] = last["age_years"].astype(int)
    last = last.sort_values(["rep", "case_no"]).drop_duplicates(["patient_initials", "sex", "age", "b", "d"])
    wk = [week(d) for d in last["d"]]
    last["y"] = [w[0] for w in wk]
    last["w"] = [w[1] for w in wk]
    band = np.where(last["age"] < 15, 0, np.where(last["age"] < 65, 1, 2))
    last["band"] = band
    idx = {n: i for i, n in enumerate(names)}
    res = np.zeros((len(names), 3))
    hist = np.zeros((len(names), 3))
    for b, y, w, a in zip(last["b"], last["y"], last["w"], last["band"]):
        if y == 2027 and 37 <= w <= 40:
            res[idx[b], a] += 1
        if y <= 2026:
            hist[idx[b], a] += 1
    n = len(names)
    # the home-address count and the placement of working-age cases at workplaces
    home = res.sum(1)
    erate = np.clip(emp / pop[:, 1], 0, 1)
    work = res[:, 1] * erate
    Wc = W[:, :n] / W[:, :n].sum(1, keepdims=True)
    placed = res[:, 0] + res[:, 2] + res[:, 1] * (1 - erate) + 0.4 * work + (0.6 * work[:, None] * Wc).sum(0)
    kappa = hist[:, 1] / np.maximum(hist[:, [0, 2]].sum(1), 1)
    ex = np.clip(res[:, 1] - kappa * res[:, [0, 2]].sum(1), 0, None)
    Wn = W[:, :n].copy()
    np.fill_diagonal(Wn, 0)
    Wn = Wn / Wn.sum(1, keepdims=True)
    A = emp[:, None] * Wn
    coef, _ = nnls(A, ex)
    inv = home - ex + coef * A.sum(0)
    hot_home = {names[i] for i in np.nonzero(home >= 40)[0]}
    hot_placed = {names[i] for i in np.nonzero(placed >= 40)[0]}
    hot_inv = {names[i] for i in np.nonzero(inv >= 40)[0]}
    gold = set(fig["call"]["hotspots"])
    ok("the call: the placement at a 0.6 share files the hotspots", hot_placed == gold, str(sorted(hot_placed)))
    ok("the call: the inversion files the hotspots", hot_inv == gold, str(sorted(hot_inv)))
    ok("the default path: the home-address count files another set", hot_home != gold and hot_home == set(fig["default_path"]),
       str(sorted(hot_home)))
    # the plan read through pypdf, with a positive control
    plan = " ".join(pg.extract_text() for pg in PdfReader(os.path.join(T, "dalisay_dengue_plan_2027_section4.pdf")).pages)
    plan = " ".join(plan.split())
    ok("pdf read: the plan's trigger (positive control)", "40 or more of those cases were infected in that barangay" in plan)
    ok("pdf read: the plan files the house index", "at least one container" in plan)
    # ask A: DMC stays stitched, DDH by date with the moved patients matched
    dmc = pd.read_csv(os.path.join(T, "dmc_dengue_inpatient_census_2027.csv"), dtype=str, keep_default_na=False)
    dmc["tin"] = pd.to_datetime(dmc["ward_in"], format="%d/%m/%Y %H:%M")
    dmc["tout"] = pd.to_datetime(dmc["ward_out"], format="%d/%m/%Y %H:%M", errors="coerce")
    dmc = dmc.sort_values(["hospital_no", "tin"])
    prev_out = dmc.groupby("hospital_no")["tout"].shift(1)
    prev_disp = dmc.groupby("hospital_no")["disposition"].shift(1)
    new_stay = ~((prev_disp == "") & (prev_out == dmc["tin"]))
    starts = dmc[new_stay]
    dmc_w = {}
    for t in starts["tin"]:
        y, w = week(t.date())
        if y == 2027 and 33 <= w <= 40:
            dmc_w[str(w)] = dmc_w.get(str(w), 0) + 1
    ok("ask A: DMC admissions recompute", {k: dmc_w.get(k, 0) for k in fig["dmc_admissions"]} == fig["dmc_admissions"])
    outs = dmc[dmc["disposition"] == "Transferred to other hospital"]
    outs = [(r.tout, int(r.age), r.sex, r.barangay.upper(), r.city_municipality.upper()) for r in outs.itertuples()]
    wb = load_workbook(os.path.join(T, "ddh_dengue_admissions_2027.xlsx"), data_only=True)
    ok("xlsx read: DDH's log has its weekly sheets (positive control)", "MW40" in wb.sheetnames)
    ddh_w, used = {}, set()
    for ws in wb.worksheets:
        for r in list(ws.iter_rows(values_only=True))[4:]:
            if not r[0]:
                continue
            t = dt.datetime.strptime(r[4] + " " + r[5], "%m/%d/%Y %I:%M %p")
            b, town = [x.strip().upper() for x in r[3].rsplit(",", 1)]
            hit = next((k for k, o in enumerate(outs) if k not in used and o[1] == int(r[1]) and o[2] == r[2]
                        and o[3] == b and o[4] == town and dt.timedelta(0) < t - o[0].to_pydatetime() <= dt.timedelta(hours=8)), None)
            if hit is not None:
                used.add(hit)
                continue
            y, w = week(t.date())
            if y == 2027 and 33 <= w <= 40:
                ddh_w[str(w)] = ddh_w.get(str(w), 0) + 1
    ok("ask A: DDH admissions recompute", {k: ddh_w.get(k, 0) for k in fig["ddh_admissions"]} == fig["ddh_admissions"])
    # ask B: houses by purok and number
    lv = pd.read_csv(os.path.join(T, "larval_survey_2027_09.csv"), dtype=str, keep_default_na=False)
    lv["pos"] = np.where(lv["team"] == "VCU", (lv["larvae"] == "Y") | (lv["pupae"] == "Y"),
                         pd.to_numeric(lv["containers_positive"], errors="coerce").fillna(0) > 0)
    h = lv.groupby(["barangay", "purok", "household_code"])["pos"].max().reset_index()
    hi = (h.groupby("barangay")["pos"].mean() * 100).round(1)
    ok("ask B: house indices recompute", all(abs(hi[b] - v) < 1e-9 for b, v in fig["house_index"].items()))
    # ask C: accepted runs, the person by initials, birth date and sex
    runs = json.load(open(os.path.join(T, "cho_lab_run_log_2027.json"), encoding="utf-8"))
    ok("json read: the run log (positive control: a rejected run)", any(r["qc"] == "Rejected" for r in runs["runs"]))
    rej = {r["run_id"] for r in runs["runs"] if r["qc"] == "Rejected"}
    lab = pd.read_csv(os.path.join(T, "cho_lab_dengue_results_2027.csv"), dtype=str, keep_default_na=False)
    lab = lab[~lab["run_id"].isin(rej)].copy()
    lab["c"] = pd.to_datetime(lab["collected_on"], format="%d/%m/%Y")
    pos = lab[lab["result"] == "Positive"].sort_values(["c", "specimen_id"])
    first = pos.drop_duplicates(["patient_initials", "date_of_birth", "sex"])
    cnt = {}
    for t, test in zip(first["c"], first["test"]):
        y, w = week(t.date())
        if y == 2027 and 35 <= w <= 40:
            k = ("NS1" if test.startswith("NS1") else "IgM", str(w))
            cnt[k] = cnt.get(k, 0) + 1
    ok("ask C: first NS1 positives recompute", {w: cnt.get(("NS1", w), 0) for w in fig["first_positive_ns1"]} == fig["first_positive_ns1"])
    ok("ask C: first IgM positives recompute", {w: cnt.get(("IgM", w), 0) for w in fig["first_positive_igm"]} == fig["first_positive_igm"])
    # the goldens
    gpath = os.path.join(TASK, "golden")
    gcsv = pd.read_csv(os.path.join(gpath, "response_figures_2027.csv"), dtype=str)
    ok("golden csv read: 40 figures", len(gcsv) == 40)
    brief = " ".join(" ".join(pg.extract_text() for pg in PdfReader(os.path.join(gpath, "dengue_hotspots_mw41_44_2027.pdf")).pages).split())
    ok("golden pdf read: the decision sentence (positive control)", "San Roque, Bagumbayan and Maligaya" in brief)
    bad = [o for o in OUT if not o[1]]
    for label, good, detail in OUT:
        print("%s  %s%s" % ("PASS" if good else "FAIL", label, ("  [%s]" % detail) if detail and not good else ""))
    print("%d of %d verifier checks pass" % (len(OUT) - len(bad), len(OUT)))
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else TASK))
