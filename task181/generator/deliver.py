"""The goldens, computed from the pack in target/: summer2027_safety_forecast.pptx (the call with the bar chart against
the tolerance line, every sector's summer 2027 forecast beside the tolerance, and the loss record it rests on) and
summer2026_sector_measures.csv (the three summer 2026 measures by sector), plus generator/figures.json, the owed figures
the write-up is drafted from."""
import csv
import io
import json
import os
import sys

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt   # noqa: E402
from pptx import Presentation     # noqa: E402
from pptx.dml.color import RGBColor   # noqa: E402
from pptx.enum.text import PP_ALIGN   # noqa: E402
from pptx.util import Cm, Pt          # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import asks as A          # noqa: E402
import models as M        # noqa: E402
import postprocess as PP  # noqa: E402
import world as W         # noqa: E402

DECK = "summer2027_safety_forecast.pptx"
CSV = "summer2026_sector_measures.csv"
NAVY, SLATE, RED, PALE, LINE, PAPER = "#203040", "#5f6b76", "#b23a2e", "#c9d2da", "#d9dee3", "#ffffff"


def compute(target):
    c = M.compute(target)
    c["asks"] = A.golden(target, rec=c["rec"], occ=c["occ"])
    S = range(8)
    smi = c["occ"][(c["occ"]["occurrence_type"] == "SMI") & c["occ"]["atm_contribution"].isin(["D", "I"])]
    c["n_losses_record"] = int(len(smi))
    return c


def r1(x):
    return "%.1f" % x


def figures(c):
    g, a = c["r"]["golden"], c["asks"]
    S = range(8)
    over = [W.NAME[W.SECTORS[s]] for s in S if g[s] > W.TOLERANCE]
    out = {
        "sectors_over_tolerance": over,
        "tolerance_per_100k_fh": "40.0",
        "forecast_rate_s27": {W.NAME[W.SECTORS[s]]: r1(g[s]) for s in S},
        "forecast_losses_s27": {W.NAME[W.SECTORS[s]]: r1(c["exp27"][s]) for s in S},
        "planning_flight_hours_s27": {W.NAME[W.SECTORS[s]]: int(round(c["ex27"][(s, "s27")]["fh"])) for s in S},
        "crossings_s27": {W.NAME[W.SECTORS[s]]: int(c["ex27"][(s, "s27")]["cross"]) for s in S},
        "crossings_record": {W.NAME[W.SECTORS[s]]: int(c["ex"][(s, "all")]["cross"]) for s in S},
        "losses_record": {W.NAME[W.SECTORS[s]]: int(c["L"][(s, "all")]) for s in S},
        "flight_hours_record": {W.NAME[W.SECTORS[s]]: int(round(c["ex"][(s, "all")]["fh"])) for s in S},
        "losses_record_total": int(sum(c["L"][(s, "all")] for s in S)),
        "crossings_record_total": int(sum(c["ex"][(s, "all")]["cross"] for s in S)),
        "losses_per_10000_crossings": "%.2f" % (c["lam"] * 1e4),
        "rate_s26_per_100k_fh": {W.NAME[W.SECTORS[s]]: r1(c["r"]["R0"][s]) for s in S},
        "crossings_s26": {W.NAME[W.SECTORS[s]]: int(c["ex"][(s, "s26")]["cross"]) for s in S},
        "flight_hours_s26": {W.NAME[W.SECTORS[s]]: int(round(c["ex"][(s, "s26")]["fh"])) for s in S},
        "rate_record_per_100k_fh": {W.NAME[W.SECTORS[s]]: r1(c["L"][(s, "all")] / c["ex"][(s, "all")]["fh"] * 1e5)
                                    for s in S},
        "losses_per_10000_crossings_by_sector": {W.NAME[W.SECTORS[s]]: "%.2f" % (c["L"][(s, "all")] /
                                                 c["ex"][(s, "all")]["cross"] * 1e4) for s in S},
        "default_path_highest": W.NAME[W.SECTORS[max(S, key=lambda s: c["r"]["R1"][s])]],
        "default_path_highest_rate": r1(max(c["r"]["R1"].values())),
        "next_highest_forecast": W.NAME[W.SECTORS[sorted(S, key=lambda s: -g[s])[1]]],
        "next_highest_forecast_rate": r1(sorted(g.values())[-2]),
        "delay_per_flight_s26_min": {W.NAME[s]: "%.2f" % a["delay"]["golden"][k] for k, s in enumerate(W.SECTORS)},
        "flights_entered_s26": {W.NAME[s]: int(a["delay"]["den"][k]) for k, s in enumerate(W.SECTORS)},
        "hours_open_alone_s26": {W.NAME[s]: r1(a["hours"][s]) for s in W.SECTORS},
        "acas_ra_encounters_s26": {W.NAME[s]: int(a["enc"][s]) for s in W.SECTORS},
        "tarvik_forecast_about": "72",
        "others_under": "30",
        "rate_basis_pairs": "10,000",
        "rate_basis_flight_hours": "100,000",
        "record_window": "1 January 2025 to 30 September 2026",
        "season": "1 June to 31 August, local time",
    }
    return out


def write_csv(path, c):
    """The summer 2026 measures, one row per sector in register order."""
    a = c["asks"]
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh, lineterminator="\r\n")
        w.writerow(["sector", "enroute_atfm_delay_per_flight_min", "flights_entering_sector", "hours_open_on_own",
                    "acas_ra_encounters"])
        for k, s in enumerate(W.SECTORS):
            w.writerow([W.NAME[s], "%.2f" % a["delay"]["golden"][k], int(a["delay"]["den"][k]), "%.1f" % a["hours"][s],
                        int(a["enc"][s])])


def write_deck(path, c, fig, chart_png):
    """Three slides for the Unit Safety Review Board: the call with the chart, the forecast by sector, and the record
    the rate rests on."""
    S = range(8)
    g = c["r"]["golden"]
    T = W.SECTORS.index("TRV")
    navy, slate, red, pale = RGBColor(0x20, 0x30, 0x40), RGBColor(0x5F, 0x6B, 0x76), RGBColor(0xB2, 0x3A, 0x2E), \
        RGBColor(0xEE, 0xF1, 0xF4)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
    blank = prs.slide_layouts[6]

    def text(slide, x, y, w, h, runs, size=12, color=slate, bold=False):
        tf = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h)).text_frame
        tf.word_wrap = True
        for k, line in enumerate(runs if isinstance(runs, list) else [runs]):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = line
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
        return tf

    def frame(slide, n, title, sub):
        text(slide, 1.6, 0.9, 30.6, 1.6, title, size=24, color=navy, bold=True)
        text(slide, 1.6, 2.5, 30.6, 1.0, sub, size=13)
        text(slide, 1.6, 17.9, 30.6, 0.7, "Holmvik ACC Safety Office  |  Unit Safety Review Board, 4 November 2026  |  "
             "Summer 2027 sector plan  |  %d" % n, size=9)

    def table(slide, rows, x, y, widths, flag_row=None, bold_last=False, right_from=1):
        h = 0.95 * len(rows)
        tb = slide.shapes.add_table(len(rows), len(rows[0]), Cm(x), Cm(y), Cm(sum(widths)), Cm(h)).table
        for j, wd in enumerate(widths):
            tb.columns[j].width = Cm(wd)
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = tb.cell(i, j)
                cell.text = val
                cell.fill.solid()
                last = bold_last and i == len(rows) - 1
                cell.fill.fore_color.rgb = navy if i == 0 else (pale if (last or i == flag_row) else
                                                                RGBColor(0xFF, 0xFF, 0xFF))
                para = cell.text_frame.paragraphs[0]
                if not para.runs:
                    para.add_run()
                run = para.runs[0]
                run.font.size = Pt(12 if i else 11)
                run.font.bold = i == 0 or last or i == flag_row
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else (red if i == flag_row else navy)
                if j >= right_from:
                    para.alignment = PP_ALIGN.RIGHT
        return tb

    s1 = prs.slides.add_slide(blank)
    frame(s1, 1, "Tarvik will run over the summer 2027 tolerance; no other sector will",
          "Forecast ATM-related losses of separation per 100,000 flight hours, planning traffic issue 1 (9 Oct 2026), "
          "against the 40.0 tolerance of HACC-SAF-PLN-03 cl. 4.1")
    s1.shapes.add_picture(chart_png, Cm(1.6), Cm(3.9), width=Cm(22.6))
    text(s1, 25.2, 4.6, 7.2, 12.0, [
        "Tarvik %.1f" % g[T],
        "",
        "The new Brekka base sends its departures and arrivals through Tarvik's trunk-airway levels: %s pairs of "
        "flights whose levels change order in the planning traffic, against %s in summer 2026, on %s flight hours "
        "against %s." % ("{:,}".format(int(c["ex27"][(T, "s27")]["cross"])), "{:,}".format(int(c["ex"][(T, "s26")]["cross"])),
                          "{:,}".format(int(round(c["ex27"][(T, "s27")]["fh"]))),
                          "{:,}".format(int(round(c["ex"][(T, "s26")]["fh"])))),
        "",
        "Every other sector is at or under %.1f." % sorted(g.values())[-2],
    ], size=12)
    s1.shapes[-1].text_frame.paragraphs[0].runs[0].font.size = Pt(22)
    s1.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True
    s1.shapes[-1].text_frame.paragraphs[0].runs[0].font.color.rgb = red

    s2 = prs.slides.add_slide(blank)
    frame(s2, 2, "Forecast by sector: only Tarvik is above 40.0",
          "Summer 2027 (1 June to 31 August, local time), planning traffic issue 1; forecast before any measure under cl. 4.2")
    rows = [["Sector", "Planning flight hours", "Pairs changing level order", "Forecast ATM-related losses",
             "Forecast rate per 100,000 FH", "Tolerance"]]
    for s in S:
        rows.append([W.NAME[W.SECTORS[s]], "{:,}".format(int(round(c["ex27"][(s, "s27")]["fh"]))),
                     "{:,}".format(int(c["ex27"][(s, "s27")]["cross"])), "%.1f" % c["exp27"][s], "%.1f" % g[s], "40.0"])
    rows.append(["All sectors", "{:,}".format(int(round(sum(c["ex27"][(s, "s27")]["fh"] for s in S)))),
                 "{:,}".format(int(sum(c["ex27"][(s, "s27")]["cross"] for s in S))), "%.1f" % sum(c["exp27"].values()),
                 "", ""])
    table(s2, rows, 1.6, 4.1, [5.4, 5.0, 5.6, 5.4, 5.6, 3.6], flag_row=1 + T, bold_last=True)
    text(s2, 1.6, 14.4, 30.6, 3.0, [
        "Forecast losses: %.2f ATM-related losses per 10,000 pairs (the record since 1 January 2025, all eight sectors) "
        "times each sector's pairs in the planning traffic. A pair is two flights in the sector at the same time whose "
        "levels change order while both are in it, from entry and exit levels taken as a straight line." % (c["lam"] * 1e4),
        "Flight hours: sector entry to exit, every flight of the planning traffic (cl. 3.5)."], size=11)

    s3 = prs.slides.add_slide(blank)
    frame(s3, 3, "Losses follow the pairs whose levels change order, not the hours flown",
          "ATM-related losses of separation (SMI, ATM contribution D or I), 1 January 2025 to 30 September 2026")
    rows = [["Sector", "ATM-related losses", "Flight hours", "Pairs changing level order", "Losses per 100,000 FH",
             "Losses per 10,000 pairs"]]
    for s in S:
        e = c["ex"][(s, "all")]
        L = c["L"][(s, "all")]
        rows.append([W.NAME[W.SECTORS[s]], str(L), "{:,}".format(int(round(e["fh"]))), "{:,}".format(int(e["cross"])),
                     "%.1f" % (L / e["fh"] * 1e5), "%.2f" % (L / e["cross"] * 1e4)])
    tl = sum(c["L"][(s, "all")] for s in S)
    tf_ = sum(c["ex"][(s, "all")]["fh"] for s in S)
    tc = sum(c["ex"][(s, "all")]["cross"] for s in S)
    rows.append(["All sectors", str(tl), "{:,}".format(int(round(tf_))), "{:,}".format(int(tc)),
                 "%.1f" % (tl / tf_ * 1e5), "%.2f" % (tl / tc * 1e4)])
    table(s3, rows, 1.6, 4.1, [5.4, 4.6, 5.0, 5.6, 5.4, 5.2], bold_last=True)
    smi = c["occ"][c["occ"]["occurrence_type"] == "SMI"]
    text(s3, 1.6, 14.4, 30.6, 3.0, [
        "Every one of the %d ATM-related losses was two aircraft in the sector together with one climbing or descending "
        "through the other's level. Losses per pair hold between %.2f and %.2f in every sector; losses per 100,000 "
        "flight hours run from %.1f to %.1f." % (tl, min(c["L"][(s, "all")] / c["ex"][(s, "all")]["cross"] * 1e4 for s in S),
                                               max(c["L"][(s, "all")] / c["ex"][(s, "all")]["cross"] * 1e4 for s in S),
                                               min(c["L"][(s, "all")] / c["ex"][(s, "all")]["fh"] * 1e5 for s in S),
                                               max(c["L"][(s, "all")] / c["ex"][(s, "all")]["fh"] * 1e5 for s in S)),
        "The %d SMI with no ATM contribution (code N, crew deviations from a correct clearance) are outside the indicator "
        "(cl. 3.4). Source: occurrence export and FDPS archive, Data and Analytics extract of 12 October 2026."
        % int((smi["atm_contribution"] == "N").sum())], size=11)
    prs.core_properties.title = "Summer 2027 safety forecast"
    prs.save(path)
    PP.normalize_ooxml(path, W.PEOPLE["safety"], W.PEOPLE["safety"], "2026-10-14T08:12:40Z", "2026-10-19T15:02:11Z",
                       title="Summer 2027 safety forecast")


def write_png(path, c):
    """The slide's chart: one bar per sector with the tolerance line; the slide carries the title and the source."""
    g = c["r"]["golden"]
    order = list(range(8))
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "svg.hashsalt": "t181",
                         "axes.edgecolor": SLATE, "axes.labelcolor": SLATE, "xtick.color": NAVY, "ytick.color": SLATE})
    fig, ax = plt.subplots(figsize=(8.0, 4.7), dpi=200)
    fig.patch.set_facecolor(PAPER)
    ax.set_facecolor(PAPER)
    names = [W.NAME[W.SECTORS[s]] for s in order]
    vals = [g[s] for s in order]
    colors = [RED if v > W.TOLERANCE else PALE for v in vals]
    ax.bar(names, vals, color=colors, width=0.62, edgecolor=PAPER, linewidth=1.5)
    ax.axhline(W.TOLERANCE, color=NAVY, linewidth=1.1, linestyle=(0, (5, 3)))
    ax.text(7.45, W.TOLERANCE + 1.2, "Tolerance 40.0", ha="right", va="bottom", fontsize=9, color=NAVY)
    for i, v in enumerate(vals):
        ax.text(i, v + 1.0, "%.1f" % v, ha="center", va="bottom", fontsize=9.5,
                color=RED if v > W.TOLERANCE else SLATE, fontweight="bold" if v > W.TOLERANCE else "normal")
    ax.set_ylabel("ATM-related losses per 100,000 flight hours")
    ax.yaxis.grid(True, color=LINE, linewidth=0.8)
    ax.set_axisbelow(True)
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    ax.tick_params(axis="x", length=0)
    ax.set_ylim(0, max(vals) * 1.15)
    fig.tight_layout()
    fig.savefig(path, format="png", facecolor=PAPER, metadata={"Software": None})
    plt.close(fig)


def main(target, golden, figures_path):
    os.makedirs(golden, exist_ok=True)
    c = compute(target)
    fig = figures(c)
    buf = io.BytesIO()
    write_png(buf, c)
    buf.seek(0)
    write_deck(os.path.join(golden, DECK), c, fig, buf)
    write_csv(os.path.join(golden, CSV), c)
    with open(figures_path, "w") as fh:
        json.dump(fig, fh, indent=1, ensure_ascii=False)
        fh.write("\n")
    return c


if __name__ == "__main__":
    task = os.path.dirname(HERE)
    main(os.path.join(task, "target"), os.path.join(task, "golden"), os.path.join(HERE, "figures.json"))
