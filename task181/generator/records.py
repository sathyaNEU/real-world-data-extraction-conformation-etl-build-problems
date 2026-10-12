"""Writers for the pack's data files, in each source system's own layout."""
import csv
import datetime as dt
import json

import numpy as np
import pyarrow as pa
import pyarrow.parquet as pq

import configs as CF
import traffic as T
import world as W

FDPS = "fdps_sector_traversals_2025-01_2026-09.parquet"
OCC = "occurrence_export_2025-01_2026-09.csv"
CODES = "occurrence_scheme_codes.csv"
PLAN = "planning_traffic_summer2027_issue1.csv"
REGISTER = "sector_register.json"
SLOTS = "nm_slot_allocations_2026-05_2026-09.csv"
REGS = "nm_regulations_2026-05_2026-09.csv"
CLOG = "sector_configuration_log_2026-05_2026-09.csv"
CTAB = "sector_configurations_current.csv"
ROSTER = "controller_roster_summer2026.xlsx"
MET = "met_cb_activity_summer2026.csv"


def _ts(sec):
    return T.at(sec)


def write_fdps(path, f, t, ids):
    order = np.lexsort((t["sec"], t["t0"]))
    fi = t["fi"][order]
    cs = f["callsign"][fi]
    tbl = pa.table({
        "ifplid": pa.array([ids[k] for k in fi], pa.string()),
        "callsign": pa.array(cs, pa.string()),
        "adep": pa.array(f["adep"][fi], pa.string()),
        "ades": pa.array(f["ades"][fi], pa.string()),
        "aircraft_type": pa.array(f["actype"][fi], pa.string()),
        "sector": pa.array(np.array(W.SECTORS)[t["sec"][order]], pa.string()),
        "entry_time_utc": pa.array((t["t0"][order] + int(T.EPOCH.timestamp())).astype("int64"), pa.timestamp("s", tz="UTC")),
        "exit_time_utc": pa.array((t["t1"][order] + int(T.EPOCH.timestamp())).astype("int64"), pa.timestamp("s", tz="UTC")),
        "entry_fl": pa.array(t["a"][order].astype("int16"), pa.int16()),
        "exit_fl": pa.array(t["b"][order].astype("int16"), pa.int16()),
    }).replace_schema_metadata({})
    pq.write_table(tbl, path, compression="zstd", row_group_size=250000, use_dictionary=["callsign", "adep", "ades",
                   "aircraft_type", "sector"], write_statistics=True)
    return tbl.num_rows


def write_plan(path, f, t):
    order = np.lexsort((t["sec"], t["t0"]))
    first = {}
    for k in order:
        fi = int(t["fi"][k])
        if fi not in first:
            first[fi] = len(first)
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh, lineterminator="\r\n")
        w.writerow(["flight_key", "cs", "dep", "dest", "ac_type", "tv", "t_in", "t_out", "fl_in", "fl_out"])
        for k in order:
            fi = int(t["fi"][k])
            w.writerow(["S27F%06d" % (first[fi] + 1), f["callsign"][fi], f["adep"][fi], f["ades"][fi], f["actype"][fi],
                        W.TV[W.SECTORS[t["sec"][k]]], _ts(t["t0"][k]).strftime("%Y-%m-%d %H:%M:%S"),
                        _ts(t["t1"][k]).strftime("%Y-%m-%d %H:%M:%S"), int(t["a"][k]), int(t["b"][k])])
    return len(order)


OCC_COLS = ["occurrence_no", "occurrence_date", "time_utc", "occurrence_type", "sector", "ac1_callsign", "ac1_level_ft",
            "ac1_phase", "ac2_callsign", "ac2_level_ft", "ac2_phase", "min_horizontal_nm", "min_vertical_ft", "severity",
            "atm_contribution", "reported_by", "report_received", "status"]


def occurrence_table(occ, seed=18163):
    rng = np.random.default_rng(seed)
    rows = []
    for o in occ:
        t = _ts(int(o["tx"]))
        rec = t + dt.timedelta(days=int(rng.choice([0, 0, 1, 1, 1, 2, 3, 5, 6])), hours=int(rng.integers(1, 9)))
        rows.append(dict(o, t=t, rec=rec))
    rows.sort(key=lambda r: (r["rec"], r["t"]))
    seq = {}
    out = []
    for r in rows:
        y = r["rec"].year
        seq[y] = seq.get(y, 0) + int(rng.integers(1, 6))      # the database numbers every unit's reports
        late = r["rec"].date() > dt.date(2026, 9, 18)
        status = "Closed" if r["type"] == "SMI" or not late or rng.random() < 0.3 else "Open"
        out.append({
            "occurrence_no": "AAN-%d-%05d" % (y, seq[y]),
            "occurrence_date": r["t"].strftime("%Y-%m-%d"),
            "time_utc": r["t"].strftime("%H:%M"),
            "occurrence_type": r["type"],
            "sector": W.NAME[W.SECTORS[r["sec"]]],
            "ac1_callsign": r["cs1"],
            "ac1_level_ft": "" if r["lv1"] is None else str(int(r["lv1"])),
            "ac1_phase": r["ph1"],
            "ac2_callsign": r["cs2"],
            "ac2_level_ft": "" if r["lv2"] is None else str(int(r["lv2"])),
            "ac2_phase": r["ph2"],
            "min_horizontal_nm": "" if r["hor"] is None else "%.1f" % r["hor"],
            "min_vertical_ft": "" if r["ver"] is None else str(int(r["ver"])),
            "severity": r["sev"],
            "atm_contribution": r["contrib"],
            "reported_by": r["rep"],
            "report_received": r["rec"].strftime("%Y-%m-%d"),
            "status": status,
        })
    out.sort(key=lambda r: (r["occurrence_date"], r["time_utc"], r["occurrence_no"]))
    return out


def write_csv(path, rows, cols):
    with open(path, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=cols, lineterminator="\n")
        w.writeheader()
        for r in rows:
            w.writerow({c: r[c] for c in cols})
    return len(rows)


CODE_ROWS = [
    ("occurrence_type", "SMI", "Separation minima infringement (loss of separation): separation between aircraft below the applicable minimum"),
    ("occurrence_type", "ACAS RA", "ACAS resolution advisory reported by a flight crew"),
    ("occurrence_type", "LB", "Level bust: deviation of 300 ft or more from the cleared level"),
    ("occurrence_type", "PLOC", "Prolonged loss of communication"),
    ("occurrence_type", "DEV", "Deviation from ATC clearance (other than a level bust)"),
    ("occurrence_type", "ATMT", "ATM technical occurrence (system or equipment)"),
    ("occurrence_type", "AI", "Airspace infringement"),
    ("occurrence_type", "WTE", "Wake turbulence encounter"),
    ("severity", "A", "Serious incident"),
    ("severity", "B", "Major incident"),
    ("severity", "C", "Significant incident"),
    ("severity", "D", "Not determined"),
    ("severity", "E", "No safety effect"),
    ("atm_contribution", "D", "Direct: an ATM event or action was in the causal chain"),
    ("atm_contribution", "I", "Indirect: ATM did not cause the occurrence but its action or inaction contributed to the outcome"),
    ("atm_contribution", "N", "No ATM contribution: no ATM event or action in the causal chain (for example a crew deviation from a correct clearance)"),
    ("atm_contribution", "U", "Not yet determined"),
    ("phase", "CLB", "Climbing"),
    ("phase", "DES", "Descending"),
    ("phase", "CRZ", "Level flight"),
    ("reported_by", "ATC", "Mandatory occurrence report by the controller"),
    ("reported_by", "STCA", "Detected by the short-term conflict alert monitoring and confirmed by the supervisor"),
    ("reported_by", "Flight crew", "Mandatory occurrence report by the operator's crew, received through the operator"),
    ("reported_by", "Engineering", "Technical occurrence report"),
]


def write_codes(path):
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["list", "code", "meaning"])
        for r in CODE_ROWS:
            w.writerow(r)
    return len(CODE_ROWS)


def write_register(path):
    reg = {
        "unit": W.UNIT, "ansp": W.ANSP, "airspace": "Holmvik FIR upper airspace, FL245 to FL660",
        "register_revision": 14, "valid_from": "2026-03-19",
        "sectors": [{"code": s, "name": W.NAME[s], "traffic_volume": W.TV[s], "lower_limit": "FL245",
                     "upper_limit": "FL660", "frequency_mhz": W.FREQ[s], "adjacent_sectors": W.ADJACENT[s]}
                    for s in W.SECTORS],
        "aerodromes_below": [{"icao": W.HUB, "name": "Holmvik", "arrival_traffic_volume": "EXHMARR"},
                             {"icao": W.BREKKA, "name": "Brekka"}, {"icao": W.SKARNES, "name": "Skarnes"}],
    }
    with open(path, "w", newline="\n") as fh:
        json.dump(reg, fh, indent=2, ensure_ascii=False)
        fh.write("\n")


def write_slots(path, rows, seed=18171):
    rng = np.random.default_rng(seed)
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["ifplid", "callsign", "adep", "ades", "eobt", "regulation_id", "slot_delay_min", "ctot"])
        for r in rows:
            eobt = _ts(r["eobt"]).replace(second=0)
            ctot = eobt + dt.timedelta(minutes=int(rng.integers(10, 19)) + r["delay"])
            w.writerow([r["ifplid"], r["callsign"], r["adep"], r["ades"], eobt.strftime("%Y-%m-%d %H:%M"), r["reg"],
                        r["delay"], ctot.strftime("%Y-%m-%d %H:%M")])
    return len(rows)


REASON = {"C": "ATC capacity", "S": "ATC staffing", "W": "Weather", "M": "Military activity", "T": "ATC equipment",
          "G": "Aerodrome capacity"}


def write_regs(path, regs):
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["regulation_id", "reg_date", "traffic_volume", "location_type", "reason", "start_utc", "end_utc",
                    "rate_per_hour"])
        for r in regs:
            w.writerow([r["id"], r["start"].strftime("%Y-%m-%d"), r["tv"], r["kind"], REASON[r["reason"]],
                        r["start"].strftime("%Y-%m-%d %H:%M"), r["end"].strftime("%Y-%m-%d %H:%M"), r["rate"]])
    return len(regs)


def write_clog(path, rows):
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["log_time_local", "configuration", "supervisor", "remarks"])
        for r in rows:
            w.writerow([r["local"].strftime("%Y-%m-%d %H:%M"), r["code"], r["sup"], r["remark"]])
    return len(rows)


def write_ctab(path):
    n = 0
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["configuration", "position", "sectors", "use"])
        for code in CF.ORDER:
            for k, g in enumerate(CF.REV9[code]):
                w.writerow([code, "E%d" % (k + 1), "+".join(g), CF.PURPOSE[code]])
                n += 1
    return n


def write_met(path, seed=18177):
    rng = np.random.default_rng(seed)
    n = 0
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["date", "sector", "cb_max_cover_pct", "cb_hours", "lightning_strikes"])
        d = dt.date(2026, 6, 1)
        while d <= dt.date(2026, 8, 31):
            storm = rng.random() < 0.28
            for s in W.SECTORS:
                cov = 0 if not storm or rng.random() < 0.45 else int(min(rng.gamma(1.6, 7), 60))
                hrs = 0.0 if cov == 0 else round(float(rng.uniform(0.5, 7.5)), 1)
                strikes = 0 if cov == 0 else int(rng.gamma(1.2, 40 * cov / 10))
                w.writerow([d.isoformat(), W.NAME[s], cov, hrs, strikes])
                n += 1
            d += dt.timedelta(days=1)
    return n


def write_roster(path, seed=18179):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    rng = np.random.default_rng(seed)
    wb = openpyxl.Workbook()
    first = True
    for month, name in ((6, "June"), (7, "July"), (8, "August")):
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name
        ws.append(["Date", "Day shift team", "Evening team", "Night team", "ATCOs day", "ATCOs evening", "ATCOs night",
                   "Supervisors", "Remarks"])
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="D9E1F2")
        d = dt.date(2026, month, 1)
        k = 0
        while d.month == month:
            teams = "ABCDEF"
            off = d.toordinal() % 6
            rem = None
            if rng.random() < 0.09:
                rem = ["%d ATCOs on refresher training" % int(rng.integers(1, 4)),
                       "1 ATCO sick, cover from team %s" % teams[(off + 3) % 6],
                       "OJT instructor on %s" % W.NAME[W.SECTORS[int(rng.integers(8))]],
                       "Overtime shift approved (%s)" % ["day", "evening", "night"][int(rng.integers(3))],
                       "Medical check, %d ATCO off roster" % 1, "Team %s short one, swap agreed" % teams[(off + 1) % 6],
                       "Simulator session for %s endorsement" % W.NAME[W.SECTORS[int(rng.integers(8))]]][
                    int(rng.integers(7))]
            ws.append([d, teams[off], teams[(off + 2) % 6], teams[(off + 4) % 6], int(rng.integers(17, 22)),
                       int(rng.integers(15, 20)), int(rng.integers(6, 9)), 2 if rng.random() < 0.8 else 1, rem])
            ws.cell(row=ws.max_row, column=1).number_format = "dd/mm/yyyy"
            d += dt.timedelta(days=1)
            k += 1
        for col, wd in zip("ABCDEFGHI", (12, 15, 13, 12, 11, 14, 12, 12, 38)):
            ws.column_dimensions[col].width = wd
        ws.freeze_panes = "A2"
    wb.save(path)
