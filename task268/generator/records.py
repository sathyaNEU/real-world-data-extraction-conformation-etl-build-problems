"""The pack's record files as each system exports them: the registry's patients, clinic visits and follow-up events,
the state's death linkage, the Tri-State exchange's death records, the record-number crosswalk, the hospital
reference, the laboratories' NT-proBNP results and the two operational files the committee pack carries. Each writer
returns its row count."""
import csv
import datetime as dt

import numpy as np

import common as K
import world as W


def _csv(path, header, rows, delim=","):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=delim, lineterminator="\r\n")
        w.writerow(header)
        w.writerows(rows)
    return len(rows)


def day(d):
    return W.DAY0 + dt.timedelta(days=int(d))


def iso(d):
    return day(d).isoformat()


def stamp(minutes):
    return (dt.datetime(2022, 1, 1) + dt.timedelta(minutes=int(minutes))).strftime("%Y-%m-%d %H:%M")


def two_year(dx):
    d = day(dx)
    return (dt.date(d.year + 2, d.month, d.day) - W.DAY0).days


def hosp_code(p, when):
    if p["hosp"] == "BRE" and p["sat"] and when >= W.SPLIT_DAY:
        return K.SATELLITE[0]
    return p["hosp"]


def spell_class(r, c, when):
    opts = K.SPELL_LEGACY[c] if when < W.MIGRATION_DAY else K.SPELL_NEW[c]
    return opts[0] if r.random() < 0.7 else opts[1]


def spell_setting(r, s, when):
    if when < W.MIGRATION_DAY:
        return {"Inpatient": "IP", "Outpatient": "OP", "Emergency department": "ED"}[s]
    if s == "Inpatient":
        return "Inpatient" if r.random() < 0.7 else "INPT"
    return {"Outpatient": "Outpatient", "Emergency department": "ED visit"}[s]


def registry_mrn(p):
    """The record number the patient was enrolled under: legacy before the June 2024 migration, new from it."""
    if p["enrol"] < W.MIGRATION_DAY and p.get("mrn_legacy"):
        return p["mrn_legacy"]
    return p["mrn_new"]


# ---------------------------------------------------------------------------------------------- patients
def write_patients(path, w):
    r = np.random.default_rng([w["seed"], 31])
    rows = []
    for p in sorted(w["patients"], key=lambda p: (p["enrol"], p["rid"])):
        if p["enrol"] > W.EXTRACT_DAY:
            continue
        county = K.COUNTY[p["hosp"]] if r.random() < 0.86 else str(r.choice(
            [K.COUNTY[h] for h in K.CODES if h != p["hosp"]] + ["Sumter", "Polk", "Manatee", "Lake"]))
        rows.append([p["rid"], registry_mrn(p), hosp_code(p, p["enrol"]), iso(p.get("dx_shown", p["dx"])), iso(p["enrol"]),
                     p["age"], p["sex"],
                     p["ef"], spell_class(r, p["c0"], p["enrol"]), spell_setting(r, p["setting"], p["enrol"]), county])
    return _csv(path, ["registry_id", "mrn", "hospital", "diagnosis_date", "enrolment_date", "age_at_diagnosis", "sex",
                       "ef_category", "nyha_at_diagnosis", "diagnosis_setting", "county_of_residence"], rows)


# ---------------------------------------------------------------------------------------------- clinic visits
def end_day(p):
    """The last day the clinic sees the patient: a move reported at a visit, the re-consent, the extract."""
    ends = [x for x in (p["north"], p["exch"], p["wd"]) if x >= 0]
    return min(ends) if ends else W.EXTRACT_DAY


def write_visits(path, w):
    r = np.random.default_rng([w["seed"], 37])
    rows = []
    for p in w["patients"]:
        if p["cls"] is None:
            continue
        stop = end_day(p)
        wt = float(r.normal(84 if p["sex"] == "M" else 72, 15))
        for k in range(1, len(p["cls"])):
            v = int(p["vday"][k])
            if not p["vatt"][k] or v > W.EXTRACT_DAY or v > stop:
                continue
            if p["wd"] >= 0 and v >= p["wd"]:
                continue
            rid = p["rid"]
            if p.get("reenrolled_as") and v >= p["reenrol_day"]:
                rid = p["reenrolled_as"]
            c = int(p["cls"][k])
            sbp = int(round(r.normal(132 - 6 * c, 14)))
            hr = int(round(r.normal(70 + 4 * c, 9)))
            loop = int(r.choice([0, 20, 40, 40, 80, 80, 120][: 3 + 2 * c] if c else [0, 0, 20, 40]))
            wt += float(r.normal(0.1 * c, 1.2))
            rows.append([rid, iso(v), hosp_code(p, v), spell_class(r, c, v), "%.1f" % wt, sbp, hr, loop])
    rows.sort(key=lambda x: (x[1], x[0]))
    return _csv(path, ["registry_id", "visit_date", "clinic", "nyha_class", "weight_kg", "systolic_bp", "heart_rate",
                       "loop_diuretic_mg_day"], rows)


# ---------------------------------------------------------------------------------------------- follow-up events
def write_events(path, w):
    r = np.random.default_rng([w["seed"], 41])
    rows = []
    for p in w["patients"]:
        if p["north"] >= 0 and p["north"] <= W.EXTRACT_DAY:
            st = str(r.choice(K.STATES_NORTH, p=np.array(K.STATE_W) / sum(K.STATE_W)))
            rows.append([p["rid"], iso(p["north"]), "REL", st])
        if p["exch"] >= 0 and p["exch"] <= W.EXTRACT_DAY:
            st = str(r.choice(K.EXCHANGE_STATES, p=(0.52, 0.18, 0.30)))
            rows.append([p["rid"], iso(p["exch"]), "REL", st])
        if p["wd"] >= 0:
            rows.append([p["rid"], iso(p["wd"]), "CNR", ""])
    rows.sort(key=lambda x: (x[1], x[0]))
    return _csv(path, ["registry_id", "event_date", "event", "destination_state"], rows)


# ---------------------------------------------------------------------------------------------- deaths
def death_source(p):
    """Which source learns of the patient's death: the Florida linkage, the exchange, or none."""
    d = p["death"]
    if d < 0 or d > W.DEATHS_DAY:
        return None
    if p["north"] >= 0 and p["north"] < d:
        return None
    if p["wd"] >= 0 and p["wd"] < d:
        return None
    if p["exch"] >= 0 and p["exch"] < d:
        return "exchange"
    return "florida"


def write_fl_deaths(path, w):
    import openpyxl
    from openpyxl.styles import Font
    r = np.random.default_rng([w["seed"], 43])
    rows = []
    for p in w["patients"]:
        if p["cls"] is None or death_source(p) != "florida":
            continue
        county = K.COUNTY[p["hosp"]] if r.random() < 0.8 else str(r.choice(["Sumter", "Polk", "Manatee", "Lake",
                                                                              "Hillsborough", "Pinellas"]))
        cause = "Cardiovascular" if r.random() < 0.62 else "Other"
        mt = "Deterministic" if r.random() < 0.9 else "Probabilistic"
        rows.append((p["mrn_new"], day(p["death"]), county, cause, mt))
    rows.sort(key=lambda x: (x[1], x[0]))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Linked deaths"
    ws.append(["MRN", "Date of death", "County of death", "Underlying cause group", "Match type"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for x in rows:
        ws.append(list(x))
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "mm/dd/yyyy"
    for col, wd in zip("ABCDE", (14, 14, 16, 22, 14)):
        ws.column_dimensions[col].width = wd
    ab = wb.create_sheet("About this run")
    for line in [["Florida death certificate linkage for the " + K.REGISTRY],
                 ["Run date", dt.date(2026, 3, 6)],
                 ["Deaths searched", "Florida resident and non-resident deaths registered in Florida, 1 January 2022 to "
                                     "31 December 2025"],
                 ["Records submitted", "Every patient in registry follow-up, by current medical record number"],
                 ["Returned", "One row per linked death"]]:
        ab.append(line)
    ab["B2"].number_format = "mm/dd/yyyy"
    ab.column_dimensions["A"].width = 20
    ab.column_dimensions["B"].width = 90
    wb.save(path)
    return len(rows)


def write_exchange(path, w):
    r = np.random.default_rng([w["seed"], 47])
    rows = []
    for p in w["patients"]:
        if p["cls"] is None or death_source(p) != "exchange":
            continue
        d = day(p["death"])
        st = str(r.choice(K.EXCHANGE_STATES, p=(0.52, 0.18, 0.30)))
        rq = "%dQ%d" % (d.year + (1 if d.month >= 10 else 0), ((d.month - 1) // 3 + 1) % 4 + 1)
        rows.append([p["rid"], st, d.strftime("%Y%m%d"), p["sex"], p["age"] + (p["death"] - p["dx"]) // 365, rq])
    rows.sort(key=lambda x: (x[2], x[0]))
    return _csv(path, ["PARTNER_ID", "STATE", "DOD", "SEX", "AGE", "RCVD_QTR"], rows, delim="|")


# ---------------------------------------------------------------------------------------------- crosswalk, hospitals
def write_crosswalk(path, w):
    r = np.random.default_rng([w["seed"], 53])
    rows = [[a, b, "2024-06-01"] for a, b in w["xwalk"]]
    order = r.permutation(len(rows))
    return _csv(path, ["legacy_mrn", "new_mrn", "mapped_on"], [rows[i] for i in order])


def write_hospitals(path):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hospitals"
    ws.append(["Code", "Hospital", "Campus", "County", "Laboratory", "Code in use from", "Code in use to"])
    for c, n, county in K.HOSPITALS:
        ws.append([c, n, "Main campus", county, c, dt.date(2015, 1, 1), None])
    ws.append([K.SATELLITE[0], K.NAME["BRE"], "North Campus", "Hernando", "BRE", K.SPLIT, None])
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        row[5].number_format = "mm/dd/yyyy"
        row[6].number_format = "mm/dd/yyyy"
    for col, wd in zip("ABCDEFG", (8, 36, 14, 14, 12, 16, 16)):
        ws.column_dimensions[col].width = wd
    nt = wb.create_sheet("Notes")
    nt.append(["The North Campus of Brenmoor Regional Medical Center opened in 2021 and used the main campus's code "
               "until its own code took effect on 1 July 2024. Its specimens go to the main campus laboratory."])
    nt.column_dimensions["A"].width = 120
    wb.save(path)
    return len(K.HOSPITALS) + 1


# ---------------------------------------------------------------------------------------------- laboratory results
def result_text(v):
    return str(W.shown(v))


def write_lab(path, w):
    rows = []
    for x in w["lab"]:
        loc = x["loc"]
        site = "BRN" if (x["lab"] == "BRE" and x["sat"] and x["col"] >= W.SPLIT_DAY * 1440) else x["lab"]
        rows.append([x["lab"], "%s%d" % ({"FEN": "F", "LIN": "L", "BRE": "B", "GUL": "G", "WRE": "W"}[x["lab"]], x["acc"]),
                     x["mrn"], stamp(x["col"]), stamp(x["rec"]), stamp(x["rel"]), "NTBNP", result_text(x["val"]), "pg/mL",
                     "F", "%s-%s" % (site, loc)])
    return _csv(path, ["performing_lab", "accession", "mrn", "collected", "received", "released", "test_code", "result",
                       "units", "result_status", "ordering_location"], rows)


# ---------------------------------------------------------------------------------------------- operational files
def write_echo(path, w):
    import openpyxl
    from openpyxl.styles import Font
    r = np.random.default_rng([w["seed"], 59])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Echo studies 2025"
    ws.append(["Month", "Hospital", "Studies", "Of which limited studies", "Reported EF under 40%", "Median days to study"])
    base = {"FEN": 610, "LIN": 470, "BRE": 520, "GUL": 690, "WRE": 540}
    for m in range(1, 13):
        for h in K.CODES:
            n = int(base[h] * (1.08 if m in (1, 2, 12) else 1.0) + r.normal(0, 25))
            ws.append([dt.date(2025, m, 1), K.NAME[h], n, int(n * r.uniform(0.08, 0.14)), int(n * r.uniform(0.17, 0.25)),
                       int(r.integers(3, 15))])
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "mmm yyyy"
    for col, wd in zip("ABCDEF", (12, 36, 10, 22, 22, 20)):
        ws.column_dimensions[col].width = wd
    wb.save(path)
    return 60


def write_clinics(path, w):
    r = np.random.default_rng([w["seed"], 61])
    rows = []
    clinicians = {"FEN": ["Dr A. Patel", "Dr M. Ochoa", "NP J. Reyes"], "LIN": ["Dr K. Brandt", "NP S. Whitley"],
                  "BRE": ["Dr R. Castellano", "Dr T. Nguyen", "NP L. Fordham"], "GUL": ["Dr D. Kessler", "Dr P. Abara",
                                                                                        "NP C. Lindqvist"],
                  "WRE": ["Dr E. Sorensen", "NP H. Mbeki"]}
    for h in K.CODES:
        for cl in clinicians[h]:
            for wd in ("Mon", "Tue", "Wed", "Thu", "Fri"):
                if r.random() < 0.55:
                    rows.append([h, cl, wd, str(r.choice(["AM", "PM"])), int(r.choice([8, 10, 12])), int(r.choice([20, 30])),
                                 "2026-01-05", "2026-12-18"])
    return _csv(path, ["hospital", "clinician", "weekday", "session", "slots", "slot_minutes", "from", "to"], rows)
