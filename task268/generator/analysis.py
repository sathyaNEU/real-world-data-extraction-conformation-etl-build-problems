"""The golden, computed from the shipped files only: homes assembled from the Rewards cards and online accounts that
pay with one bank card (each till payment matched to its Faircourt row on terminal, time and amount; each online
order on its order number), each rival's buying homes and overlap with the Roastery, every rung and hazard of the
ladder, and the two asks."""
import bisect
import collections
import csv
import datetime as dt
import itertools
import json
import os
import re

import common as K

PACK_FORMATS = ("Ground", "Beans", "Capsules", "Instant")


def norm_brand(s):
    s = re.sub(r"[^A-Z ]", "", s.upper().replace("'", ""))
    words = [w for w in s.split() if w not in ("COFFEE", "ROASTERS", "CAFFE", "KAFFE")]
    key = " ".join(words)
    table = {"CORCALLYS ROASTERY": "OL", "HALLORANS": "HAL", "BELLVALE": "BEL", "NORDHAVN": "NOR",
             "ORO ROSSO": "ROS", "CAFETTO": "CAF", "WESTQUAY": "WES", "KINVARA HILLS": "KIN"}
    return table[key]


class UF:
    def __init__(self):
        self.p = {}

    def f(self, x):
        p = self.p
        p.setdefault(x, x)
        r = x
        while p[r] != r:
            r = p[r]
        while p[x] != r:
            p[x], x = r, p[x]
        return r

    def u(self, a, b):
        ra, rb = self.f(a), self.f(b)
        if ra != rb:
            if ra < rb:
                ra, rb = rb, ra
            self.p[ra] = rb


class Golden:
    def __init__(self, target):
        self.t = target
        p = lambda k: os.path.join(target, K.FILES[k])
        self._master(p("master"))
        self._register(p("register"))
        self._till(p("till"))
        self._settlement(p("settle"))
        self._online(p("online"))
        self.labels_path, self.deals_path = p("labels"), p("deals")

    # ------------------------------------------------------------------ loading
    def _master(self, path):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True).active
        self.brand_of, self.fmt_of, self.supplier_of, self.price_of = {}, {}, {}, {}
        for r in list(ws.iter_rows(values_only=True))[1:]:
            plu, desc, brand, fmt, g, price, supplier, online = r
            if fmt in PACK_FORMATS:
                self.brand_of[str(plu)] = norm_brand(brand)
                self.fmt_of[str(plu)] = fmt
                self.supplier_of[str(plu)] = supplier
                self.price_of[str(plu)] = price
        self.spellings = collections.defaultdict(set)
        ws = load_workbook(path, read_only=True).active
        for r in list(ws.iter_rows(values_only=True))[1:]:
            if r[3] in PACK_FORMATS:
                self.spellings[norm_brand(r[2])].add(r[2])

    def _register(self, path):
        self.store_cards, self.succ, self.reg = set(), {}, {}
        with open(path, encoding="utf-8") as f:
            for r in csv.DictReader(f):
                self.reg[r["card_no"]] = r
                if r["card_type"] == "STORE":
                    self.store_cards.add(r["card_no"])
                if r["replaced_by"]:
                    self.succ[r["card_no"]] = r["replaced_by"]

    def fold(self, card):
        seen = 0
        while card in self.succ and seen < 10:
            card = self.succ[card]
            seen += 1
        return card

    def _till(self, path):
        self.tx = {}
        with open(path, encoding="utf-8") as f:
            for r in csv.DictReader(f):
                key = (r["store"], int(r["till"]), r["receipt"], r["txn_time"])
                t = self.tx.get(key)
                if t is None:
                    t = self.tx[key] = dict(store=r["store"], till=int(r["till"]), receipt=r["receipt"],
                                            when=dt.datetime.strptime(r["txn_time"], "%Y-%m-%d %H:%M:%S"),
                                            card=r["rewards_card"], tender=r["tender"],
                                            card_amount=float(r["card_amount"]), total=float(r["txn_total"]),
                                            lines=[])
                t["lines"].append(dict(plu=r["plu"], desc=r["description"], qty=int(r["qty"]), value=float(r["value"]),
                                       promo=r["promo"], note=r["note"]))
        self.txs = sorted(self.tx.values(), key=lambda t: (t["when"], t["store"], t["till"], t["receipt"]))

    def _settlement(self, path):
        import pyarrow.parquet as pq
        tb = pq.read_table(path).to_pydict()
        n = len(tb["amount"])
        self.rows = [{k: tb[k][i] for k in tb} for i in range(n)]
        for r in self.rows:
            r["utc"] = r["txn_time_utc"].replace(tzinfo=None)

    def _online(self, path):
        self.online = json.load(open(path, encoding="utf-8"))

    # ------------------------------------------------------------------ the terminal schedule (merchant agreement)
    def tid(self, store, till, when, current_only=False):
        if (store, till) == ("LSG", 2):
            if current_only:
                return "41023212"
            return "41023202" if when < K.SWAP_DATE else "41023212"
        return [tid for s, tl, tid, a, b in K.TERMINALS if s == store and tl == till][0]

    # ------------------------------------------------------------------ till payments matched to Faircourt
    def match_day(self, current_only=False):
        """Each card-paid receipt matched on terminal, local day and amount alone."""
        idx = collections.defaultdict(list)
        for i, r in enumerate(self.rows):
            if r["terminal_id"] in (K.WEB_TID, K.COUNTER_TID) or r["txn_type"] != 'SALE':
                continue
            idx[(r["terminal_id"], r["utc"].date(), round(r["amount"], 2))].append(i)
        out = {}
        for key, t in self.tx.items():
            if t["card_amount"] <= 0:
                continue
            tid = self.tid(t["store"], t["till"], t["when"], current_only)
            hits = idx.get((tid, t["when"].date(), round(t["card_amount"], 2)), [])
            if len(hits) == 1:
                out[key] = self.rows[hits[0]]["card_ref"]
        return out

    def match(self, utc_raw=False, current_only=False, window=120):
        """Each till receipt paid by card, matched to its settlement row: same terminal, same type, same amount, the
        Faircourt time (UTC) up to `window` seconds before the receipt's Irish time. Returns {receipt key: card_ref}."""
        idx = collections.defaultdict(list)
        for i, r in enumerate(self.rows):
            if r["terminal_id"] in (K.WEB_TID, K.COUNTER_TID):
                continue
            idx[(r["terminal_id"], r["txn_type"], round(r["amount"], 2))].append((r["utc"], i))
        for v in idx.values():
            v.sort()
        out, ambiguous, missing = {}, 0, 0
        for key, t in self.tx.items():
            if t["card_amount"] <= 0:
                continue
            typ = 'SALE'
            tid = self.tid(t["store"], t["till"], t["when"], current_only)
            local = t["when"]
            target = local if utc_raw else K.to_utc(local)
            lst = idx.get((tid, typ, round(abs(t["card_amount"]), 2)), [])
            lo = bisect.bisect_left(lst, (target - dt.timedelta(seconds=window), -1))
            hi = bisect.bisect_right(lst, (target + dt.timedelta(seconds=5), 10 ** 9))
            hits = lst[lo:hi]
            if len(hits) == 1:
                out[key] = self.rows[hits[0][1]]["card_ref"]
            elif len(hits) > 1:
                ambiguous += 1
            else:
                missing += 1
        self.match_stats = (len(out), ambiguous, missing)
        return out

    def web_refs(self):
        return {r["merchant_ref"]: r["card_ref"] for r in self.rows
                if r["terminal_id"] == K.WEB_TID and r["txn_type"] == 'SALE' and r["merchant_ref"]}

    # ------------------------------------------------------------------ purchases
    def store_purchases(self, refunds_removed=False):
        """{receipt key: set of brands} for 2026 pack lines bought (qty > 0) on a till receipt."""
        out = {}
        for key, t in self.tx.items():
            if t["when"].year != K.YEAR:
                continue
            bs = set()
            for ln in t["lines"]:
                if ln["plu"] in self.brand_of and ln["qty"] > 0:
                    bs.add(self.brand_of[ln["plu"]])
            if bs:
                out[key] = bs
        if refunds_removed:
            # a pack refunded later on the same card: drop that brand from the card's receipt that carried it
            ref = collections.defaultdict(list)
            for key, t in self.tx.items():
                for ln in t["lines"]:
                    if ln["plu"] in self.brand_of and ln["qty"] < 0 and t["card"]:
                        ref[t["card"]].append((t["when"], self.brand_of[ln["plu"]]))
            for card, items in ref.items():
                for when, b in items:
                    cands = [k for k, t in self.tx.items() if t["card"] == card and t["when"] < when and k in out
                             and b in out[k]]
                    if cands:
                        last = max(cands, key=lambda k: self.tx[k]["when"])
                        out[last] = out[last] - {b}
        return out

    def online_purchases(self, basis="delivered", versions="final", cancelled=False):
        """[(account, order no, brands)] for the online orders that count."""
        out = []
        for o in self.online["orders"]:
            if o["status"] == "cancelled" and not cancelled:
                continue
            if basis == "delivered":
                if not o["delivered_on"] or not o["delivered_on"].startswith(str(K.YEAR)):
                    continue
            else:
                if not o["placed"].startswith(str(K.YEAR)):
                    continue
            vs = o["versions"] if versions == "all" else o["versions"][-1:]
            bs = set()
            for v in vs:
                for ln in v["coffee_lines"]:
                    if ln["sku"] in self.brand_of:
                        bs.add(self.brand_of[ln["sku"]])
            if bs:
                out.append((o["account"], o["order"], bs))
        return out

    # ------------------------------------------------------------------ buying units
    def units(self, kind="homes", fold=True, online=True, accounts_apart=False, hub=False, utc_raw=False,
              thr=1, current_only=False, drop_unlinked_accounts=False, refunds_removed=False, basis="delivered",
              versions="final", cancelled=False, house_lines=False, match_mode="time"):
        """Each buying unit's set of brands. kind: 'card' (each card, or member when fold), 'homes'."""
        sp = self.store_purchases(refunds_removed)
        on = self.online_purchases(basis, versions, cancelled) if (online or accounts_apart) else []
        key_of = (lambda c: self.fold(c)) if fold else (lambda c: c)
        if kind == "card":
            g = collections.defaultdict(set)
            for k, bs in sp.items():
                c = self.tx[k]["card"]
                if c and c not in self.store_cards:
                    g[key_of(c)] |= bs
            for acc, _, bs in on:
                g[("acc", acc)] |= bs
            return list(g.values())
        m = self.match_day(current_only) if match_mode == "day" else self.match(utc_raw=utc_raw, current_only=current_only)
        uf = UF()
        cnt = collections.Counter()
        for k, ref in m.items():
            c = self.tx[k]["card"]
            if not c or (c in self.store_cards and not hub):
                continue
            if self.tx[k]["card_amount"] < 0:
                continue
            cnt[(ref, key_of(c))] += 1
        web = self.web_refs()
        if online and not accounts_apart:
            for acc, no, bs in self.online_purchases("delivered", "final", False):
                if no in web:
                    cnt[(web[no], ("acc", acc))] += 1
        by_ref = collections.defaultdict(list)
        for (ref, node), n in sorted(cnt.items(), key=lambda x: str(x[0])):
            if n >= thr:
                by_ref[ref].append(node)
        for ref, nodes in by_ref.items():
            for nd in nodes[1:]:
                uf.u(("n", str(nodes[0])), ("n", str(nd)))
        g = collections.defaultdict(set)
        linked = set()
        for ref, nodes in by_ref.items():
            if len(nodes) > 1:
                linked |= {str(x) for x in nodes}
        for k, bs in sp.items():
            c = self.tx[k]["card"]
            if c and (c not in self.store_cards or hub):
                g[uf.f(("n", str(key_of(c))))] |= bs
            elif house_lines and k in m and m[k] in by_ref:
                # a store-card or anonymous receipt paid with a bank card a home's cards use
                g[uf.f(("n", str(by_ref[m[k]][0])))] |= bs
        for acc, _, bs in on:
            node = ("n", str(("acc", acc)))
            if accounts_apart:
                g[("apart", acc)] |= bs
            elif drop_unlinked_accounts and str(("acc", acc)) not in linked:
                continue
            else:
                g[uf.f(node)] |= bs
        self.last_uf = uf
        return list(g.values())

    @staticmethod
    def overlap(units):
        out = {}
        for b in K.RIVALS:
            buyers = [s for s in units if b in s]
            both = sum(1 for s in buyers if "OL" in s)
            out[b] = (len(buyers), both, 100.0 * both / len(buyers) if buyers else float("nan"))
        out["OL"] = (sum(1 for s in units if "OL" in s), None, None)
        return out

    @staticmethod
    def pick(ov):
        s = sorted(K.RIVALS, key=lambda b: ov[b][2])
        return s[0], s[1], ov[s[1]][2] / ov[s[0]][2]

    def baskets(self):
        sp = self.store_purchases()
        units = list(sp.values()) + [bs for _, _, bs in self.online_purchases()]
        return units

    def reverse(self, units):
        ol = [s for s in units if "OL" in s]
        return {b: (len(ol), sum(1 for s in ol if b in s), 100.0 * sum(1 for s in ol if b in s) / len(ol))
                for b in K.RIVALS}

    def ladder(self):
        L = collections.OrderedDict()
        L["R0 baskets"] = self.overlap(self.baskets())
        L["R1 cards"] = self.overlap(self.units("card", fold=False, online=False))
        L["R1 members"] = self.overlap(self.units("card", fold=True, online=False))
        L["R1 members and accounts"] = self.overlap(self.units("card", fold=True, online=True))
        L["R2 homes from store cards"] = self.overlap(self.units("homes", online=False))
        L["R2 accounts kept apart"] = self.overlap(self.units("homes", accounts_apart=True))
        L["R3 homes"] = self.overlap(self.units("homes"))
        return L

    def hazards(self):
        H = collections.OrderedDict()
        H["store cards kept"] = self.overlap(self.units("homes", hub=True))
        H["clock read as Irish time"] = self.overlap(self.units("homes", utc_raw=True))
        H["matched on the day"] = self.overlap(self.units("homes", match_mode="day"))
        H["two shared payments"] = self.overlap(self.units("homes", thr=2))
        H["three shared payments"] = self.overlap(self.units("homes", thr=3))
        H["replaced cards kept apart"] = self.overlap(self.units("homes", fold=False))
        H["current terminal only"] = self.overlap(self.units("homes", current_only=True))
        H["unlinked accounts dropped"] = self.overlap(self.units("homes", drop_unlinked_accounts=True))
        H["refunded packs removed"] = self.overlap(self.units("homes", refunds_removed=True))
        H["online by order date"] = self.overlap(self.units("homes", basis="placed"))
        H["every version counted"] = self.overlap(self.units("homes", versions="all"))
        H["cancelled orders counted"] = self.overlap(self.units("homes", cancelled=True))
        H["store-card lines put to homes"] = self.overlap(self.units("homes", house_lines=True))
        return H

    # ------------------------------------------------------------------ ask A: the counter
    def counter_labels(self, fix_grams=True, fix_clock=True):
        labs = []
        with open(self.labels_path, encoding="utf-8") as f:
            for r in csv.DictReader(f):
                w = float(r["net_weight"])
                if fix_grams and r["scale"] == "S2":
                    w = w / 1000.0
                labs.append(dict(id=r["label"], scale=r["scale"], no=int(r["label"].split("-")[1]),
                                 shown=dt.datetime.strptime(r["printed"], "%Y-%m-%d %H:%M:%S"), roaster=r["roaster"],
                                 coffee=r["coffee"], kg=w, price=round(float(r["label_price"]), 2)))
        for lb in labs:
            lb["when"] = lb["shown"]
        if fix_clock:
            self._fix_clock(labs)
        return labs

    def _fix_clock(self, labs):
        """A scale whose label numbers run on while its printed times jump back: the block between the jump and the
        return is re-dated by the offset that lines its labels up with the counter's card payments."""
        pays = self.counter_payments()
        for sc in ("S1", "S2"):
            seq = sorted((lb for lb in labs if lb["scale"] == sc), key=lambda x: x["no"])
            start = end = None
            for i in range(1, len(seq)):
                if start is None and seq[i]["shown"] < seq[i - 1]["shown"] - dt.timedelta(days=2):
                    start = i
                elif start is not None and seq[i]["shown"] > seq[i - 1]["shown"] + dt.timedelta(days=20):
                    end = i
                    break
            if start is None:
                continue
            block = seq[start:end]
            by_amt = collections.defaultdict(list)
            for p in pays:
                by_amt[p["amount"]].append(p["when"])
            diffs = []
            for lb in block:
                for pw in by_amt.get(lb["price"], []):
                    d = (pw - lb["shown"]).total_seconds()
                    if 86400 < d < 366 * 86400:
                        diffs.append(d)
            diffs.sort()
            # the densest ten-minute bin of differences, then its median
            best, bi = 0, 0
            j = 0
            for i in range(len(diffs)):
                while diffs[i] - diffs[j] > 600:
                    j += 1
                if i - j + 1 > best:
                    best, bi = i - j + 1, j
            core = diffs[bi:bi + best]
            off = core[len(core) // 2] - 100 if core else 0
            for lb in block:
                lb["when"] = lb["shown"] + dt.timedelta(seconds=off)
            self.clock_block = (sc, block[0]["id"], block[-1]["id"], off)

    def counter_payments(self, utc_raw=False):
        out = []
        for r in self.rows:
            if r["terminal_id"] == K.COUNTER_TID and r["txn_type"] == 'SALE':
                out.append(dict(when=r["utc"] if utc_raw else K.to_local(r["utc"]), amount=round(r["amount"], 2)))
        out.sort(key=lambda p: p["when"])
        return out

    def counter_till_lines(self):
        out = []
        for t in self.txs:
            for ln in t["lines"]:
                if ln["plu"] == K.COUNTER_PLU:
                    out.append(dict(when=t["when"], value=round(ln["value"], 2)))
        return out

    def counter_sold(self, labs, card=True, till=True, utc_raw=False):
        acc = [0.0]
        prices = [a[1] for a in K.ACCESSORIES]
        for a in prices:
            acc.append(a)
        for a, b in itertools.combinations_with_replacement(prices, 2):
            acc.append(a + b)
        accset = set(round(x, 2) for x in acc)
        times = sorted(labs, key=lambda x: x["when"])
        tkeys = [lb["when"] for lb in times]
        used = set()
        sold = set()
        self.counter_ambiguous = 0
        if card:
            for p in self.counter_payments(utc_raw):
                lo = bisect.bisect_left(tkeys, p["when"] - dt.timedelta(minutes=25))
                hi = bisect.bisect_right(tkeys, p["when"] + dt.timedelta(minutes=3))
                cands = [lb for lb in times[lo:hi] if lb["id"] not in used]
                fits = []
                for k in (1, 2, 3):
                    for comb in itertools.combinations(cands, k):
                        rest = round(p["amount"] - sum(lb["price"] for lb in comb), 2)
                        if rest in accset:
                            fits.append(comb)
                if len(fits) >= 1:
                    if len(fits) > 1:
                        self.counter_ambiguous += 1
                    for lb in fits[0]:
                        used.add(lb["id"])
                        sold.add(lb["id"])
        if till:
            for ln in self.counter_till_lines():
                lo = bisect.bisect_left(tkeys, ln["when"] - dt.timedelta(minutes=150))
                hi = bisect.bisect_right(tkeys, ln["when"] + dt.timedelta(minutes=1))
                cands = [lb for lb in times[lo:hi] if lb["id"] not in used and lb["price"] == ln["value"]]
                if cands:
                    if len(cands) > 1:
                        self.counter_ambiguous += 1
                    lb = cands[-1]
                    used.add(lb["id"])
                    sold.add(lb["id"])
        return sold

    def ask_a(self, mode="golden"):
        """kg sold by roaster and quarter. Modes: golden; every label; labels sold but the clock left; card payments
        only; over-cleaned (a label followed within three minutes by another of the same coffee on its scale dropped)."""
        if mode == "every":
            labs = self.counter_labels(fix_clock=False)
            sold = {lb["id"] for lb in labs}
        elif mode == "every_fixed":
            labs = self.counter_labels(fix_clock=True)
            sold = {lb["id"] for lb in labs}
        elif mode == "noclock":
            labs = self.counter_labels(fix_clock=False)
            sold = self.counter_sold(labs)
        elif mode == "cardonly":
            labs = self.counter_labels()
            sold = self.counter_sold(labs, till=False)
        elif mode == "overclean":
            labs = self.counter_labels(fix_clock=True)
            seq = sorted(labs, key=lambda x: (x["scale"], x["no"]))
            drop = set()
            for a, b in zip(seq, seq[1:]):
                if a["scale"] == b["scale"] and a["coffee"] == b["coffee"] and \
                        0 <= (b["when"] - a["when"]).total_seconds() <= 180:
                    drop.add(a["id"])
            sold = {lb["id"] for lb in labs} - drop
        elif mode == "grams":
            labs = self.counter_labels(fix_grams=False)
            sold = self.counter_sold(labs)
        else:
            labs = self.counter_labels()
            sold = self.counter_sold(labs)
        out = {(r, q): 0.0 for r in [x[1] for x in K.ROASTERS] for q in (1, 2, 3, 4)}
        for lb in labs:
            if lb["id"] in sold and lb["when"].year == K.YEAR:
                out[(lb["roaster"], K.quarter(lb["when"].date()))] += lb["kg"]
        return out

    # ------------------------------------------------------------------ ask B: multibuy funding
    def deal_log(self):
        from openpyxl import load_workbook
        ws = load_workbook(self.deals_path, read_only=True)["Deal log 2026"]
        deals = collections.OrderedDict()
        for r in list(ws.iter_rows(values_only=True))[1:]:
            deal, supplier, brand, plus, mech, start, end, fund, ver, agreed, eff = r
            d = deals.setdefault(deal, dict(supplier=supplier, brand=brand, plus=set(plus.split(", ")),
                                            start=start.date(), end=end.date(), price=float(mech.split()[-1]),
                                            rates=[]))
            d["rates"].append((eff.date(), float(fund)))
        for d in deals.values():
            d["rates"].sort()
        return deals

    @staticmethod
    def rate(d, day):
        r = None
        for eff, f in d["rates"]:
            if eff <= day:
                r = f
        return r

    def ask_b(self, mode="golden"):
        """Funding by supplier and quarter. Modes: golden; savings (saving lines counted, reversals netted on their
        own receipt, adjustments ignored); count_lines (every saving line counted, reversals as savings); adj_own_date
        (adjustments dated by their own receipt); adj_one (each adjustment line one multibuy); latest_rate (the
        latest version's rate on every sale)."""
        deals = self.deal_log()
        byrc = {(t["store"], t["till"], t["receipt"]): t for t in self.txs}
        out = collections.defaultdict(float)
        for t in self.txs:
            day = t["when"].date()
            for ln in t["lines"]:
                if ln["desc"] == "MULTIBUY SAVING" and ln["promo"] in deals:
                    d = deals[ln["promo"]]
                    sale_day = day
                    q = ln["qty"]
                    if q < 0 and mode != "count_lines":
                        # a reversal: cancel the multibuy on the receipt the refund cites
                        orig = self._cited(t, byrc, "refund")
                        if orig is not None and mode != "savings":
                            sale_day = orig["when"].date()
                    if mode == "count_lines":
                        q = 1
                    if not (d["start"] <= sale_day <= d["end"]):
                        continue
                    rate = d["rates"][-1][1] if mode == "latest_rate" else self.rate(d, sale_day)
                    out[(d["supplier"], K.quarter(sale_day))] += q * rate
                elif ln["plu"] == K.ADJ_PLU and mode not in ("savings", "count_lines"):
                    orig = self._cited(t, byrc, "MB missed")
                    if orig is None:
                        continue
                    oday = orig["when"].date()
                    left = round(-ln["value"], 2)
                    for did, d in deals.items():
                        if not (d["start"] <= oday <= d["end"]):
                            continue
                        packs = [x for x in orig["lines"] if x["plu"] in d["plus"] and x["qty"] > 0]
                        n = sum(x["qty"] for x in packs)
                        given = sum(1 for x in orig["lines"] if x["promo"] == did and x["qty"] > 0)
                        missed = n // 2 - given
                        if missed < 1 or left <= 0:
                            continue
                        unit = [self.price_of[x["plu"]] for x in packs for _ in range(x["qty"])]
                        save = round(unit[0] + unit[1] - d["price"], 2)
                        if mode == "adj_one":
                            pairs = 1
                        else:
                            pairs = min(missed, int(round(left / save)))
                        left = round(left - pairs * save, 2)
                        sday = day if mode == "adj_own_date" else oday
                        rate = d["rates"][-1][1] if mode == "latest_rate" else self.rate(d, oday)
                        out[(d["supplier"], K.quarter(sday))] += pairs * rate
        sup = sorted(set(d["supplier"] for d in deals.values()))
        return {(s, q): round(out[(s, q)], 2) for s in sup for q in (1, 2, 3, 4)}

    @staticmethod
    def _cited(t, byrc, word):
        for ln in t["lines"]:
            m = re.search(r"rcpt (\w{3}) (\d) (\d{6}) (\d\d)/(\d\d)", ln["note"] or "")
            if m and (word in ln["note"]):
                return byrc.get((m.group(1), int(m.group(2)), m.group(3)))
        return None

    # ------------------------------------------------------------------ the golden in one place
    def run(self):
        R = {}
        homes = self.units("homes")
        ov = self.overlap(homes)
        call, runner, margin = self.pick(ov)
        R["homes_n"] = len(homes)
        R["overlap"] = ov
        R["call"], R["runner"], R["margin"] = call, runner, margin
        R["gap"] = ov[runner][2] - ov[call][2]
        R["match_stats"] = self.match_stats
        dflt = self.overlap(self.units("card", fold=True, online=True))
        d = self.pick(dflt)
        R["default"] = (d[0], dflt[d[0]][2])
        R["default_overlap"] = dflt
        R["ask_a"] = self.ask_a()
        R["clock_block"] = self.clock_block
        R["counter_ambiguous"] = self.counter_ambiguous
        R["ask_b"] = self.ask_b()
        return R
