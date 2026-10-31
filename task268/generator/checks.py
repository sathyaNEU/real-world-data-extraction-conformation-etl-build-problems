"""The generator's assertions on the cut pack and the seeded world: input gates, the ladder and its margins, every
semantic fork row, the three rules (R1 to R3c) on the clean truth, each Layer 1 item's neutrality, the separation of
the asks from the call, both asks against the truth and their stops, the distractors, the realism sweep and the
documents' positive controls.

    python3 task268/generator/checks.py            (run by build.py)"""
import collections
import csv
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import analysis as A   # noqa: E402
import common as K     # noqa: E402

TASK = os.path.dirname(HERE)


def text_of(path):
    """Text of a shipped document through the same extractors verify.py uses."""
    if path.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(path).pages)
    if path.endswith(".docx"):
        import docx
        return "\n".join(p.text for p in docx.Document(path).paragraphs)
    if path.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True)
        return "\n".join(" ".join(str(c) for c in r if c is not None) for ws in wb for r in ws.iter_rows(values_only=True))
    if path.endswith(".parquet"):
        import pyarrow.parquet as pq
        t = pq.read_table(path)
        return " ".join(t.column_names)
    return open(path, encoding="utf-8").read()


def same(a, b):
    return all(a[x][:2] == b[x][:2] for x in K.RIVALS)


def rounded(ov):
    return {b: (ov[b][0], ov[b][1], int(round(ov[b][2]))) for b in K.RIVALS}


def run(task=TASK, world=None):
    T = os.path.join(task, "target")
    rows = []

    def ok(name, cond, detail=""):
        rows.append((name, bool(cond), detail))

    G = A.Golden(T)
    R = G.run()
    ov, call, runner = R["overlap"], R["call"], R["runner"]
    files = sorted(os.listdir(T))
    # ---------------------------------------------------------------- input gates
    ok("gate files", len(files) >= 10, "%d files" % len(files))
    fmts = sorted(set(f.rsplit(".", 1)[1] for f in files))
    ok("gate formats", len(fmts) >= 3, ", ".join(fmts))
    n_till = sum(1 for _ in open(os.path.join(T, K.FILES["till"]), encoding="utf-8")) - 1
    ok("gate large file", n_till >= 25000 and len(G.rows) >= 25000, "%d till lines, %d settlement rows" % (n_till, len(G.rows)))
    meta = json.load(open(os.path.join(task, "metadata.json")))
    dis = meta.get("distractor_files", [])
    ok("gate distractors", len(dis) >= 2 and all(os.path.exists(os.path.join(task, d)) for d in dis), ", ".join(dis))
    bad = [f for f in files if re.search(r"readme|dictionary|manifest|notes|field|schema|distractor", f, re.I)]
    sheets = []
    from openpyxl import load_workbook
    for f in files:
        if f.endswith(".xlsx"):
            for ws in load_workbook(os.path.join(T, f), read_only=True).worksheets:
                if re.search(r"readme|notes|dictionary|about|how to|fields", ws.title, re.I):
                    sheets.append(f + ":" + ws.title)
    ok("gate no file or sheet describes the data", not bad and not sheets, ", ".join(bad + sheets) or "none")
    # ---------------------------------------------------------------- the ladder
    L = G.ladder()
    picks = {k: G.pick(v) for k, v in L.items()}
    for k, want, floor in (("R1 cards", "NOR", 1.15), ("R1 members", "NOR", 1.15), ("R1 members and accounts", "NOR", 1.15),
                           ("R2 homes from store cards", "WES", 1.15), ("R2 accounts kept apart", "WES", 1.15),
                           ("R3 homes", "ROS", 1.2)):
        c, r, m = picks[k]
        ok("rung %s names %s" % (k, K.BRAND_NAME[want]), c == want and m >= floor,
           "%s %.1f, next %s %.1f, x%.2f" % (c, L[k][c][2], r, L[k][r][2], m))
    ok("rung R0 baskets not the golden", picks["R0 baskets"][0] != "ROS", picks["R0 baskets"][0])
    third = sorted(K.RIVALS, key=lambda b: ov[b][2])[2]
    ok("runner-up separated from third", ov[third][2] - ov[runner][2] >= 1.0,
       "%s %.2f, %s %.2f" % (runner, ov[runner][2], third, ov[third][2]))
    edges = {b: abs(ov[b][2] - int(ov[b][2]) - 0.5) for b in K.RIVALS}
    ok("every share mid-bin", min(edges.values()) >= 0.03 and edges[call] >= 0.15 and edges[runner] >= 0.15,
       "closest %.2f; call %.2f, runner-up %.2f" % (min(edges.values()), edges[call], edges[runner]))
    ok("gap rounds alike from unrounded and rounded shares",
       int(round(R["gap"])) == int(round(ov[runner][2])) - int(round(ov[call][2])) and abs(R["gap"] % 1 - 0.5) >= 0.1,
       "%.2f" % R["gap"])
    ok("match unique", R["match_stats"][1] == 0 and R["match_stats"][2] == 0, str(R["match_stats"]))
    # reflex sweep
    online_only = G.overlap(_accounts(G))
    ok("reflex online accounts alone not the golden", G.pick(online_only)[0] != "ROS", G.pick(online_only)[0])
    heavy = G.overlap(_heavy(G))
    ok("reflex heavy cards not the golden", G.pick(heavy)[0] != "ROS", G.pick(heavy)[0])
    rev = G.reverse(G.units("homes"))
    rmin = min(K.RIVALS, key=lambda b: rev[b][2])
    ok("SF14 reverse duplication names another brand", rmin != "ROS", rmin)
    # ---------------------------------------------------------------- semantic fork rows
    H = G.hazards()
    ok("SF1 a card as the home files another brand", picks["R1 cards"][0] != call, picks["R1 cards"][0])
    ok("SF2 store-card and anonymous lines put to homes keep the call", G.pick(H["store-card lines put to homes"])[0] == call)
    ok("SF3 the counter's PLUs map to no brand", K.COUNTER_PLU not in G.brand_of and K.ACCESSORY_PLU not in G.brand_of)
    split = _split(G)
    ok("SF4 spellings kept apart leave the call and its components", split[0] == call and split[1] == rounded(ov)[call]
       and split[2] == rounded(ov)[runner], str(split))
    ok("SF5 online orders by order date leave every figure", same(H["online by order date"], ov))
    ok("SF6 every version and cancelled orders leave every figure",
       same(H["every version counted"], ov) and same(H["cancelled orders counted"], ov))
    sk = H["store cards kept"]
    ok("SF7 store cards kept leave the call and its components", G.pick(sk)[0] == call and
       rounded(sk)[call][2] == rounded(ov)[call][2] and rounded(sk)[runner][2] == rounded(ov)[runner][2],
       "%s %.2f" % (G.pick(sk)[0], sk[call][2]))
    ok("SF8 replaced cards kept apart leave every figure", same(H["replaced cards kept apart"], ov))
    ok("SF9 two or three shared payments leave every figure",
       same(H["two shared payments"], ov) and same(H["three shared payments"], ov))
    ok("SF10 the settlement clock: matched on the day leaves every figure", same(H["matched on the day"], ov))
    ok("SF10 the settlement clock read as Irish time keeps the call", G.pick(H["clock read as Irish time"])[0] == call)
    ok("SF11 the current terminal only leaves the call", G.pick(H["current terminal only"])[0] == call)
    ok("SF12 unlinked online accounts dropped keep the call", G.pick(H["unlinked accounts dropped"])[0] == call)
    ok("SF13 refunded packs removed leave every figure", same(H["refunded packs removed"], ov))
    ok("SF15 the call's share 0.15 or more from its rounding edge", edges[call] >= 0.15, "%.2f" % edges[call])
    ok("SF16 every rung margin 1.15 or more", all(p[2] >= 1.15 for k, p in picks.items() if k != "R0 baskets"))
    # ---------------------------------------------------------------- the three rules on the clean truth
    from world import World
    w = world or World(K.SEED)
    truth = _truth(w)
    ok("R1 no shipped column holds a home or household", not any(
        re.search(r"^(home|household|hh|family)(_?(id|no|ref|key))?$|home_?id|household_?id", h, re.I)
        for h in _headers(T)), "%d headers read" % len(_headers(T)))
    ok("R1 the decisive computation reads row-level receipts and settlement rows", R["match_stats"][0] > 25000)
    ok("R2 brand spellings: the call and its components hold when kept apart", split[0] == call)
    ok("R2 lost-card replacements: every figure holds unfolded", same(H["replaced cards kept apart"], ov))
    ok("R2 store cards: the call holds when kept", G.pick(sk)[0] == call)
    ok("R2 settlement in UTC: every figure holds matched on the day", same(H["matched on the day"], ov))
    ok("R2 the replaced Lisgarvan terminal: the call holds on the current terminal",
       G.pick(H["current terminal only"])[0] == call)
    ok("R2 amended and cancelled online orders: every figure holds", same(H["every version counted"], ov))
    ok("R2 refunds: every figure holds", same(H["refunded packs removed"], ov))
    ok("R3a the easy route on the clean truth files Nordhavn", truth["members"] == "NOR", truth["members"])
    ok("R3b only the unit changed (true homes) files Oro Rosso", truth["homes"] == "ROS", truth["homes"])
    ok("R3b homes from store cards alone stay wrong", picks["R2 homes from store cards"][0] not in ("ROS", "NOR"))
    ok("R3c no duplicated till line", _no_dupes(T))
    # ---------------------------------------------------------------- asks
    ta = truth["ask_a"]
    ok("ask A golden equals the truth", all(abs(R["ask_a"][k] - ta[k]) < 1e-6 for k in ta),
       "; ".join("%s %.3f/%.3f" % (k, R["ask_a"][k], ta[k]) for k in ta if abs(R["ask_a"][k] - ta[k]) >= 1e-6))
    ok("ask A matching unambiguous", R["counter_ambiguous"] == 0, str(R["counter_ambiguous"]))
    for mode in ("every", "every_fixed", "noclock", "cardonly", "overclean", "grams"):
        m = G.ask_a(mode)
        miss = sum(1 for k in ta if abs(round(m[k], 1) - round(ta[k], 1)) >= 0.05)
        ok("ask A stop %s misses" % mode, miss >= 6, "%d of 12" % miss)
    ok("ask A figures mid-bin", min(abs((v * 10) % 1 - 0.5) for v in ta.values()) >= 0.02,
       "%.3f" % min(abs((v * 10) % 1 - 0.5) for v in ta.values()))
    tb = truth["ask_b"]
    gb = R["ask_b"]
    ok("ask B golden equals the truth", all(abs(gb[k] - tb.get(k, 0.0)) < 0.005 for k in gb))
    cr = G.ask_b("savings")
    ok("ask B the cracker misses every cell", all(round(cr[k]) != round(gb[k]) for k in gb),
       "%d kept" % sum(1 for k in gb if round(cr[k]) == round(gb[k])))
    for mode in ("count_lines", "adj_own_date", "adj_one", "latest_rate"):
        m = G.ask_b(mode)
        miss = sum(1 for k in gb if round(m[k]) != round(gb[k]))
        ok("ask B stop %s misses" % mode, miss >= 3, "%d of 20" % miss)
    ok("ask B figures mid-bin", min(abs(v % 1 - 0.5) for v in gb.values()) >= 0.015,
       "%.2f" % min(abs(v % 1 - 0.5) for v in gb.values()))
    # ---------------------------------------------------------------- separation
    dev = 0
    for t in G.txs:
        for ln in t["lines"]:
            if ln["plu"] in (K.COUNTER_PLU, K.ACCESSORY_PLU, K.ADJ_PLU) or ln["desc"] == "MULTIBUY SAVING":
                dev += ln["plu"] in G.brand_of
    ok("separation: no device or hazard row is a brand pack line", dev == 0, str(dev))
    # ---------------------------------------------------------------- distractors
    src = open(os.path.join(HERE, "analysis.py"), encoding="utf-8").read()
    ok("distractors unread by the golden", all(K.FILES[k] not in src and ('"%s"' % k) not in src for k in K.DISTRACTORS))
    # ---------------------------------------------------------------- realism sweep
    ok("realism no arithmetic id column", _no_arith(T))
    counts = [ov[b][0] for b in K.RIVALS]
    ok("realism distinct buying-home counts", len(set(counts)) == len(counts))
    ok("realism no round headline", all(c % 100 for c in counts) and ov[call][1] % 10)
    late = [r for r in G.rows if r["utc"].year != K.YEAR]
    ok("realism settlement inside 2026", not late, "%d rows outside" % len(late))
    ok("realism till lines inside 2026", all(t["when"].year == K.YEAR for t in G.txs))
    em = [f for f in files if f.endswith((".csv", ".json")) and ("—" in open(os.path.join(T, f), encoding="utf-8").read()
                                                               or "–" in open(os.path.join(T, f), encoding="utf-8").read())]
    ok("realism no em or en dash in text files", not em, ", ".join(em))
    # ---------------------------------------------------------------- documents: positive controls
    brief = text_of(os.path.join(T, K.FILES["brief"]))
    ok("doc brief carries the estimand", "if anyone in it bought a pack of that brand" in brief and
       "An online order counts on the day it was delivered" in brief)
    terms = text_of(os.path.join(T, K.FILES["terms"]))
    ok("doc terms carry the personal card", "Each card is issued to one member" in terms.replace("\n", " "))
    agr = text_of(os.path.join(T, K.FILES["agreement"]))
    ok("doc agreement carries UTC and the replaced terminal", "UTC" in agr and "41023212" in agr and "41023202" in agr
       and K.COUNTER_TID in agr)
    board = text_of(os.path.join(T, K.FILES["board"]))
    ok("doc price board carries accessories and the cash route", all("%.2f" % p in board for _, p in K.ACCESSORIES)
       and "checkout" in board)
    for k in ("brief", "terms", "agreement", "board"):
        t = text_of(os.path.join(T, K.FILES[k]))
        ok("doc %s names no payment link" % k, not re.search(r"household|same (bank )?card|shared card|link", t, re.I))
    return rows


def _accounts(G):
    acc = collections.defaultdict(set)
    for a, no, bs in G.online_purchases():
        acc[a] |= bs
    return list(acc.values())


def _heavy(G):
    sp = G.store_purchases()
    per = collections.defaultdict(list)
    for k, bs in sp.items():
        c = G.tx[k]["card"]
        if c and c not in G.store_cards:
            per[G.fold(c)].append(bs)
    return [set().union(*v) for v in per.values() if len(v) >= 8]


def _split(G):
    """The golden with each master spelling a brand of its own."""
    from openpyxl import load_workbook
    ws = load_workbook(os.path.join(G.t, K.FILES["master"]), read_only=True).active
    spell = {}
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if r[3] in A.PACK_FORMATS:
            spell[str(r[0])] = r[2]
    saved = dict(G.brand_of)
    canon = {"OL": "Corcally's Roastery", "ROS": "Oro Rosso", "NOR": "Nordhavn"}
    for plu, b in saved.items():
        if spell[plu] != canon.get(b, spell[plu]) or (b not in canon and spell[plu] != K.BRAND_NAME[b]):
            G.brand_of[plu] = b + "*" if b not in canon else b
    try:
        units = G.units("homes")
    finally:
        G.brand_of = saved
    ov = G.overlap(units)
    extra = set().union(*units) - set(K.B)
    allb = list(K.RIVALS) + sorted(extra)
    shares = {}
    for b in allb:
        buyers = [s for s in units if b in s]
        if buyers:
            shares[b] = 100.0 * sum(1 for s in buyers if "OL" in s) / len(buyers)
    low = min(shares, key=shares.get)
    return low, (ov["ROS"][0], ov["ROS"][1], int(round(ov["ROS"][2]))), (ov["NOR"][0], ov["NOR"][1], int(round(ov["NOR"][2])))


def _truth(w):
    """The clean truth: the easy route and the true homes on the world before any file, and both asks."""
    out = {}
    fold = {c: m["card"] for m in w.members for c in m["cards"]}
    mem, home = collections.defaultdict(set), collections.defaultdict(set)
    for t in w.tx:
        if t["kind"] == 'sale' and t["card"] == 'MEMBER':
            bs = {ln["brand"] for ln in t["lines"] if ln["qty"] > 0}
            mem[fold[t["card_no"]]] |= bs
            home[t["home"]] |= bs
    for o in w.orders:
        if o["status"] == 'delivered' and o["delivery"].year == K.YEAR:
            bs = {K.SKU[c["plu"]][1] for c in o["versions"][-1]["coffee"]}
            mem[("acc", o["account"])] |= bs
            home[o["home"]] |= bs
    out["members"] = A.Golden.pick(A.Golden.overlap(list(mem.values())))[0]
    out["homes"] = A.Golden.pick(A.Golden.overlap(list(home.values())))[0]
    ta = {(n, q): 0.0 for _, n, _ in K.ROASTERS for q in (1, 2, 3, 4)}
    for lb in w.labels:
        if lb["sold"]:
            ta[(K.ROASTER_NAME[lb["roaster"]], K.quarter(lb["when"].date()))] += lb["kg"]
    out["ask_a"] = ta

    def rate(d, day):
        r = d[6]
        for eff, f in d[7]:
            if eff <= day:
                r = f
        return r
    tb = collections.defaultdict(float)
    for t in w.tx:
        if t["kind"] == 'sale':
            for s in t["saving"]:
                d = K.DEAL[s["deal"]]
                tb[(K.SUPPLIER[d[1]], K.quarter(t["when"].date()))] += rate(d, t["when"].date())
        elif t["kind"] == 'refund':
            for s in t["saving"]:
                d = K.DEAL[s["deal"]]
                o = t["orig"]
                tb[(K.SUPPLIER[d[1]], K.quarter(o["when"].date()))] -= rate(d, o["when"].date())
    for a in w.adjust:
        for c in a["claims"]:
            d = K.DEAL[c["deal"]]
            o = c["tx"]
            tb[(K.SUPPLIER[d[1]], K.quarter(o["when"].date()))] += c["pairs"] * rate(d, o["when"].date())
    out["ask_b"] = {k: round(v, 2) for k, v in tb.items()}
    return out


def _headers(T):
    hs = []
    for f in os.listdir(T):
        p = os.path.join(T, f)
        if f.endswith(".csv"):
            hs += next(csv.reader(open(p, encoding="utf-8")))
        elif f.endswith(".parquet"):
            import pyarrow.parquet as pq
            hs += pq.read_schema(p).names
        elif f.endswith(".json"):
            d = json.load(open(p, encoding="utf-8"))
            hs += list(d.keys()) + list(d["orders"][0].keys())
    return hs


def _no_dupes(T):
    seen = set()
    with open(os.path.join(T, K.FILES["till"]), encoding="utf-8") as f:
        for line in f:
            if line in seen:
                return False
            seen.add(line)
    return True


def _no_arith(T):
    for f, col in ((K.FILES["register"], "card_no"), (K.FILES["register"], "member_no"),
                   (K.FILES["till"], "receipt")):
        vals = []
        with open(os.path.join(T, f), encoding="utf-8") as fh:
            for r in csv.DictReader(fh):
                v = re.sub(r"\D", "", r[col])
                if v:
                    vals.append(int(v))
        vals = sorted(set(vals))
        if len(vals) > 3 and len(set(b - a for a, b in zip(vals, vals[1:]))) == 1:
            return False
    return True


if __name__ == "__main__":
    rs = run()
    bad = [r for r in rs if not r[1]]
    for r in rs:
        print("%s  %-70s %s" % ("PASS" if r[1] else "FAIL", r[0], r[2]))
    print("checks: %d assertions, %d failed" % (len(rs), len(bad)))
    sys.exit(1 if bad else 0)
