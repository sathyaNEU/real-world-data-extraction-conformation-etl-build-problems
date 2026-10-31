"""Write the evidence pack from the seeded world: each system's file in its own format and idiom."""
import csv
import datetime as dt
import io
import json
import os

import numpy as np

import common as K
import docs
import world as W


def _f2(x):
    return "%.2f" % x


def till_rows(w):
    """The coffee category's till lines as the warehouse extract holds them: one row a line, the receipt's own
    fields on every line."""
    out = []
    for t in sorted(w.tx, key=lambda x: (x["when"], x["store"], x["till"], x["receipt"])):
        tender = "CASH" if t["card_amount"] == 0 else ("CARD+CASH" if abs(t["card_amount"]) < abs(t["total"]) - 0.001
                                                       else "CARD")
        lines = []
        if t["kind"] in ('sale', 'refund'):
            merged = {}
            for ln in t["lines"]:
                k = ln["plu"]
                if k not in merged:
                    merged[k] = [0, 0.0]
                merged[k][0] += ln["qty"]
                merged[k][1] += ln["value"]
            note = ""
            if t["kind"] == 'refund':
                o = t["orig"]
                note = "refund rcpt %s %d %06d %s" % (o["store"], o["till"], o["receipt"], o["when"].strftime("%d/%m"))
            for plu, (q, v) in merged.items():
                lines.append((plu, K.SKU[plu][3], q, K.r2(v), "", note))
            for s in t["saving"]:
                q = 1 if s["value"] < 0 else -1
                lines.append(("", "MULTIBUY SAVING", q, K.r2(s["value"]), s["deal"], ""))
        elif t["kind"] == 'counter':
            for lb in t["counter"]:
                lines.append((K.COUNTER_PLU, "COUNTER COFFEE LOOSE", 1, lb["price"], "", ""))
        elif t["kind"] == 'adjust':
            o = t["orig"]
            note = "MB missed rcpt %s %d %06d %s" % (o["store"], o["till"], o["receipt"], o["when"].strftime("%d/%m"))
            lines.append((K.ADJ_PLU, "MULTIBUY ADJ", 1, K.r2(t["value"]), "", note))
        for i, (plu, desc, q, v, promo, note) in enumerate(lines, 1):
            out.append([t["store"], t["till"], "%06d" % t["receipt"], t["when"].strftime("%Y-%m-%d %H:%M:%S"), i, plu,
                        desc, q, _f2(v), promo, t["card_no"], tender, _f2(t["card_amount"]), _f2(t["total"]), note])
    return out


TILL_HEAD = ["store", "till", "receipt", "txn_time", "line", "plu", "description", "qty", "value", "promo",
             "rewards_card", "tender", "card_amount", "txn_total", "note"]


def write_csv(path, head, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        wr = csv.writer(f, lineterminator="\r\n")
        wr.writerow(head)
        wr.writerows(rows)


def online_json(w):
    orders = []
    for o in w.orders:
        if o["delivery"] > dt.date(2027, 1, 7):
            continue
        vs = []
        t = o["placed"]
        for i, v in enumerate(o["versions"], 1):
            if i > 1:
                t = t + dt.timedelta(hours=int(3 + 7 * i), minutes=17 * i)
            lines = [dict(sku=c["plu"], item=K.SKU[c["plu"]][4], qty=c["qty"], price=K.SKU[c["plu"]][7])
                     for c in v["coffee"]]
            total = K.r2(sum(K.SKU[c["plu"]][7] * c["qty"] for c in v["coffee"]) + v["other"])
            vs.append(dict(version=i, saved=t.strftime("%Y-%m-%dT%H:%M:%S"), coffee_lines=lines, basket_total=total))
        orders.append(dict(order=o["no"], account=o["account"], placed=o["placed"].strftime("%Y-%m-%dT%H:%M:%S"),
                           slot_date=o["delivery"].strftime("%Y-%m-%d"), status=o["status"],
                           delivered_on=o["delivery"].strftime("%Y-%m-%d") if o["status"] == 'delivered' else None,
                           versions=vs))
    doc = dict(export="order history, coffee lines", shop="Corcally's online shop (Drumsallagh fulfilment)",
               exported_at=K.EXTRACT.strftime("%Y-%m-%dT%H:%M:%S"), placed_from=K.ONLINE_FROM.isoformat(),
               timezone="Europe/Dublin", orders=orders)
    return doc


def register_rows(w):
    rows = []
    for no in sorted(w.cards):
        c = w.cards[no]
        mno = "" if c["type"] == "STORE" else w.member_no[c["member"]]
        rows.append([no, mno, c["type"], c["issued"].isoformat(), c["status"],
                     c["status_date"].isoformat() if c["status_date"] else "", c["replaced_by"] or "", c["store"],
                     c.get("age", "")])
    return rows


REG_HEAD = ["card_no", "member_no", "card_type", "issued", "status", "status_date", "replaced_by", "home_store",
            "age_band"]


def write_settlement(path, w):
    import pyarrow as pa
    import pyarrow.parquet as pq
    rows = w.settle
    cols = {
        "merchant_id": pa.array([r["merchant_id"] for r in rows], pa.string()),
        "terminal_id": pa.array([r["terminal_id"] for r in rows], pa.string()),
        "txn_time_utc": pa.array([r["txn_time_utc"] for r in rows], pa.timestamp("s", tz="UTC")),
        "txn_type": pa.array([r["txn_type"] for r in rows], pa.string()),
        "amount": pa.array([r["amount"] for r in rows], pa.float64()),
        "currency": pa.array([r["currency"] for r in rows], pa.string()),
        "card_ref": pa.array([r["card_ref"] for r in rows], pa.string()),
        "scheme": pa.array([r["scheme"] for r in rows], pa.string()),
        "card_product": pa.array([r["card_product"] for r in rows], pa.string()),
        "card_region": pa.array([r["card_region"] for r in rows], pa.string()),
        "entry_mode": pa.array([r["entry_mode"] for r in rows], pa.string()),
        "auth_code": pa.array([r["auth_code"] for r in rows], pa.string()),
        "merchant_ref": pa.array([r["merchant_ref"] or None for r in rows], pa.string()),
        "settle_date": pa.array([r["settle_date"] for r in rows], pa.date32()),
        "interchange": pa.array([r["interchange"] for r in rows], pa.float64()),
        "scheme_fee": pa.array([r["scheme_fee"] for r in rows], pa.float64()),
        "service_charge": pa.array([r["service_charge"] for r in rows], pa.float64()),
        "net_settled": pa.array([r["net_settled"] for r in rows], pa.float64()),
    }
    table = pa.table(cols)
    table = table.replace_schema_metadata({"source": "Faircourt Merchant Services, merchant settlement detail",
                                           "merchant": "Corcally & Sons Limited", "period": "2026-01-01/2026-12-31",
                                           "time_basis": "UTC"})
    pq.write_table(table, path, compression="zstd", row_group_size=65536, use_dictionary=True,
                   write_statistics=True)


def label_rows(w):
    rows = []
    for lb in sorted(w.labels, key=lambda x: (x["shown"], x["scale"], x["no"])):
        weight = ("%.3f" % lb["kg"]) if lb["scale"] == "S1" else ("%d" % round(lb["kg"] * 1000))
        rows.append([lb["scale"] + "-%06d" % lb["no"], lb["scale"], lb["shown"].strftime("%Y-%m-%d %H:%M:%S"),
                     K.ROASTER_NAME[lb["roaster"]], lb["coffee"], weight, "%.2f" % lb["ppk"], _f2(lb["price"])])
    return rows


LABEL_HEAD = ["label", "scale", "printed", "roaster", "coffee", "net_weight", "price_per_kg", "label_price"]


def write_master(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Coffee range"
    ws.append(["PLU", "Description", "Brand", "Sub-category", "Pack (g)", "Shelf price (EUR)", "Supplier",
               "Listed online"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for s in sorted(K.SKUS, key=lambda x: (x[5], x[2], x[0])):
        ws.append([s[0], s[4], s[2], s[5], s[6], s[7], K.SUPPLIER[s[1]], s[8]])
    ws.append([K.COUNTER_PLU, "Coffee counter loose beans (weighed, price embedded)", None, "Counter", None, None,
               "Various (guest roasters)", "N"])
    ws.append([K.ACCESSORY_PLU, "Coffee counter accessories (price embedded)", None, "Counter", None, None,
               "Various", "N"])
    for col, wd in zip("ABCDEFGH", (16, 52, 24, 13, 9, 16, 30, 13)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A2"
    wb.save(path)


def write_deals(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Deal log 2026"
    ws.append(["Deal", "Supplier", "Brand", "PLUs", "Mechanic", "Start", "End",
               "Funding per multibuy sold (EUR)", "Version", "Agreed", "Effective from"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for d in K.DEALS:
        deal_id, b, plus, price, start, end, fund, revs = d
        agreed = start - dt.timedelta(days=24 + int(deal_id[-2:]) % 9)
        ws.append([deal_id, K.SUPPLIER[b], K.BRAND_NAME[b], ", ".join(plus), "Any 2 for EUR %.2f" % price, start, end,
                   fund, 1, agreed, start])
        for i, (eff, f2) in enumerate(revs, 2):
            ws.append([deal_id, K.SUPPLIER[b], K.BRAND_NAME[b], ", ".join(plus), "Any 2 for EUR %.2f" % price, start, end,
                       f2, i, eff - dt.timedelta(days=5), eff])
    for row in ws.iter_rows(min_row=2):
        for c in (row[5], row[6], row[9], row[10]):
            c.number_format = "dd/mm/yyyy"
        row[7].number_format = "0.00"
    for col, wd in zip("ABCDEFGHIJK", (10, 30, 14, 34, 20, 12, 12, 16, 9, 12, 14)):
        ws.column_dimensions[col].width = wd
    wb.save(path)


def census_rows(w):
    rng = np.random.default_rng(K.SEED + 77)
    rows = []
    eds = [("Drumsallagh Urban", 18), ("Drumsallagh Rural", 9), ("Lisgarvan", 7), ("Kilbrannagh", 6),
           ("Cloonmore", 4), ("Ardnageeha", 4), ("Rathkeel", 3)]
    n = 0
    for ed, k in eds:
        for i in range(k):
            n += 1
            tot = int(rng.integers(80, 260))
            p = rng.dirichlet([2.6, 3.0, 1.7, 1.6, 1.0])
            c = np.floor(p * tot).astype(int)
            c[1] += tot - c.sum()
            persons = int(c[0] + 2 * c[1] + 3 * c[2] + 4 * c[3] + 5.6 * c[4])
            rows.append(["A%09d" % (127000000 + n * 37), ed, tot] + [int(x) for x in c] + [persons])
    return rows


CENSUS_HEAD = ["small_area", "electoral_division", "households", "hh_1_person", "hh_2_persons", "hh_3_persons",
               "hh_4_persons", "hh_5_plus_persons", "persons_in_private_households"]


def write_tracker(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    rng = np.random.default_rng(K.SEED + 91)
    wb = Workbook()
    ws = wb.active
    ws.title = "Brand health"
    ws.append(["Coffee brand tracker, Corcally's shopper panel, Midlands. Online survey of adults who buy coffee."])
    ws.append([])
    ws.append(["Wave", "Brand", "Prompted awareness %", "Ever bought %", "Bought in last 3 months %",
               "Main brand at home %", "Would consider %", "Base (n)"])
    for c in ws[3]:
        c.font = Font(bold=True)
    base = {"May 2026": 612, "Nov 2026": 588}
    lvl = {"OL": (71, 44, 29, 17, 52), "HAL": (93, 78, 41, 24, 61), "BEL": (84, 57, 22, 9, 37),
           "NOR": (66, 31, 15, 8, 26), "ROS": (58, 27, 13, 7, 30), "CAF": (91, 74, 33, 18, 44),
           "WES": (34, 12, 6, 2, 21), "KIN": (49, 22, 11, 5, 33)}
    for wave in ("May 2026", "Nov 2026"):
        for b in K.B:
            a = [int(round(v + rng.normal(0, 1.4))) for v in lvl[b]]
            ws.append([wave, K.BRAND_NAME[b]] + a + [base[wave]])
    ws.append([])
    ws.append(["Fieldwork by Lisnagree Research for Corcally's. Percentages of all respondents."])
    for col, wd in zip("ABCDEFGH", (11, 22, 21, 15, 24, 21, 17, 10)):
        ws.column_dimensions[col].width = wd
    wb.save(path)


def write_pack(out, w=None):
    w = w or W.World()
    p = lambda k: os.path.join(out, K.FILES[k])
    rows = till_rows(w)
    write_csv(p("till"), TILL_HEAD, rows)
    with open(p("online"), "w", encoding="utf-8") as f:
        json.dump(online_json(w), f, indent=1, ensure_ascii=False)
        f.write("\n")
    write_csv(p("register"), REG_HEAD, register_rows(w))
    write_settlement(p("settle"), w)
    write_csv(p("labels"), LABEL_HEAD, label_rows(w))
    write_master(p("master"))
    write_deals(p("deals"))
    write_csv(p("census"), CENSUS_HEAD, census_rows(w))
    write_tracker(p("tracker"))
    docs.write_agreement(p("agreement"))
    docs.write_terms(p("terms"))
    docs.write_board(p("board"))
    docs.write_brief(p("brief"))
    n = {"till": len(rows), "settle": len(w.settle), "labels": len(w.labels), "register": len(w.cards)}
    return w, n
