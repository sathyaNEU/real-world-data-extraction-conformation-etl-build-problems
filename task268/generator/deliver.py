"""The golden deliverables: roastery_sampling_2027.pptx (the call, the overlap chart, the rival table, how the homes
were built, the counter's loose coffee and the multibuy funding owed) and rival_home_overlap_2026.csv."""
import csv

import common as K

RNAME = K.BRAND_NAME


def pct(x):
    return int(round(x))


def figures(R):
    ov = R["overlap"]
    call, runner = R["call"], R["runner"]
    F = {"call": RNAME[call], "runner_up": RNAME[runner],
         "call_share": pct(ov[call][2]), "runner_share": pct(ov[runner][2]),
         "gap_points": int(round(R["gap"])),
         "call_homes": ov[call][0], "call_both": ov[call][1],
         "runner_homes": ov[runner][0], "runner_both": ov[runner][1],
         "homes_total": R["homes_n"], "receipts_matched": R["match_stats"][0],
         "default_call": RNAME[R["default"][0]], "default_share": pct(R["default"][1]),
         "rivals": {RNAME[b]: {"homes": ov[b][0], "both": ov[b][1], "share": pct(ov[b][2])} for b in K.RIVALS},
         "ask_a": {"%s Q%d" % k: round(v, 1) for k, v in sorted(R["ask_a"].items())},
         "ask_b": {"%s Q%d" % k: int(round(v)) for k, v in sorted(R["ask_b"].items())},
         "clock_block": list(R["clock_block"][1:3]),
         "counter_tid": K.COUNTER_TID}
    return F


def order(R):
    ov = R["overlap"]
    return sorted(K.RIVALS, key=lambda b: ov[b][2])


def write_csv(path, R):
    ov = R["overlap"]
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\r\n")
        w.writerow(["rival_brand", "homes_buying_2026", "homes_also_buying_roastery", "share_also_roastery_pct"])
        for b in order(R):
            w.writerow([RNAME[b], ov[b][0], ov[b][1], pct(ov[b][2])])


def write_deck(path, R, F):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    from pptx.util import Emu, Pt

    BROWN, GREY, INK, PALE = RGBColor(0x6B, 0x3E, 0x1F), RGBColor(0x9A, 0xA4, 0xAB), RGBColor(0x22, 0x26, 0x2A), \
        RGBColor(0xC9, 0xB8, 0xA6)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]
    ov = R["overlap"]
    call, runner = R["call"], R["runner"]

    def text(slide, x, y, w, h, s, size=14, bold=False, color=INK):
        tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, ln in enumerate(s.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = ln
            p.font.size, p.font.bold, p.font.color.rgb, p.font.name = Pt(size), bold, color, "Calibri"
        return tb

    def foot(slide, n):
        text(slide, 457200, 6380000, 9800000, 300000, "Corcally's customer insight. Rewards cards in the three stores and "
             "online shop accounts, 1 January to 31 December 2026.", 9, color=GREY)
        text(slide, 11200000, 6380000, 600000, 300000, str(n), 9, color=GREY)

    def table(slide, x, y, w, rows, widths, size=12):
        from pptx.enum.text import PP_ALIGN
        shp = slide.shapes.add_table(len(rows), len(rows[0]), Emu(x), Emu(y), Emu(w), Emu(380000 * len(rows)))
        tb = shp.table
        tb.first_row, tb.horz_banding = True, False
        for j, wd in enumerate(widths):
            tb.columns[j].width = Emu(wd)
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                c = tb.cell(i, j)
                c.text = str(v)
                c.fill.solid()
                c.fill.fore_color.rgb = BROWN if i == 0 else (RGBColor(0xF6, 0xF1, 0xEB) if i % 2 == 0 else
                                                              RGBColor(0xFF, 0xFF, 0xFF))
                p = c.text_frame.paragraphs[0]
                p.font.size, p.font.name = Pt(size), "Calibri"
                p.font.bold = i == 0
                p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else INK
                p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
        return tb

    # 1. the call
    s = prs.slides.add_slide(blank)
    text(s, 457200, 1500000, 11200000, 1100000, "%s's homes get the 2027 Roastery sampling" % RNAME[call]
         if not RNAME[call].endswith("s") else "%s' homes get the 2027 Roastery sampling" % RNAME[call], 34, True, BROWN)
    text(s, 457200, 2800000, 11000000, 1600000,
         "Of the %s homes that bought %s in 2026, %s also bought the Roastery: %d%%, the lowest of the rivals.\n"
         "%s is next at %d%%, %d points higher." % (
             "{:,}".format(ov[call][0]), RNAME[call], "{:,}".format(ov[call][1]), F["call_share"], RNAME[runner],
             F["runner_share"], F["gap_points"]), 18)
    foot(s, 1)
    # 2. the chart
    s = prs.slides.add_slide(blank)
    text(s, 457200, 300000, 11200000, 700000, "%s's buyers are the least likely to have the Roastery already"
         % RNAME[call], 26, True, BROWN)
    text(s, 457200, 950000, 11200000, 400000, "Share of each rival's 2026 buying homes that also bought the Roastery",
         14, color=GREY)
    cd = CategoryChartData()
    rows = order(R)
    cd.categories = [RNAME[b] for b in rows]
    cd.add_series("Also bought the Roastery (%)", [pct(ov[b][2]) for b in rows])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Emu(457200), Emu(1450000), Emu(10800000), Emu(4700000), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = '0"%"', False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(12)
    ser = plot.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = PALE
    pt = ser.points[0]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = BROWN
    cat = ch.category_axis
    cat.reverse_order = True
    cat.tick_labels.font.size = Pt(12)
    va = ch.value_axis
    va.maximum_scale, va.minimum_scale = 100, 0
    va.has_major_gridlines = False
    va.tick_labels.font.size = Pt(10)
    foot(s, 2)
    # 3. the table
    s = prs.slides.add_slide(blank)
    text(s, 457200, 300000, 11200000, 700000, "%s sits %d points under the next rival" % (RNAME[call], F["gap_points"]),
         26, True, BROWN)
    rows_t = [["Rival", "Homes buying in 2026", "Also bought the Roastery", "Share"]]
    for b in rows:
        rows_t.append([RNAME[b], "{:,}".format(ov[b][0]), "{:,}".format(ov[b][1]), "%d%%" % pct(ov[b][2])])
    table(s, 457200, 1200000, 9000000, rows_t, (3000000, 2000000, 2400000, 1600000))
    text(s, 457200, 4500000, 11000000, 900000, "Runner-up %s at %d%%; the gap to %s is %d points." % (
        RNAME[runner], F["runner_share"], RNAME[call], F["gap_points"]), 14)
    text(s, 457200, 5300000, 11000000, 700000, "Homes: Rewards cards and online shop accounts that paid with the same "
         "bank card in Faircourt's 2026 settlement. Receipts without a Rewards card and the stores' own cards are outside "
         "the brief.", 10, color=GREY)
    foot(s, 3)
    # 4. how the homes were built
    s = prs.slides.add_slide(blank)
    text(s, 457200, 300000, 11200000, 700000, "Counting homes, not cards", 26, True, BROWN)
    text(s, 457200, 1250000, 11000000, 4200000,
         "A Rewards card belongs to one member (Rewards terms 2.1), so two people in one home carry two cards, and the "
         "online shop keeps its own accounts.\n"
         "Every 2026 coffee receipt paid by card was matched to its Faircourt settlement row on terminal, amount and "
         "time (Faircourt's UTC time against the till's Irish time); every delivered online order to its row by order "
         "number.\n"
         "Cards and accounts that paid with the same bank card were joined into one home. The stores' own cards, swiped "
         "for shoppers without one, were left out.\n"
         "Card by card the picture is different: %s's capsule buyers are mostly a home's second card, and another "
         "member of the home buys the Roastery." % RNAME[R["default"][0]], 15)
    foot(s, 4)
    # 5. the counter
    s = prs.slides.add_slide(blank)
    tot = {name: sum(R["ask_a"][(name, q)] for q in (1, 2, 3, 4)) for _, name, _ in K.ROASTERS}
    top = max(tot, key=tot.get)
    text(s, 457200, 300000, 11200000, 700000, "%s sold the most loose coffee at the Drumsallagh counter" % top, 26,
         True, BROWN)
    text(s, 457200, 950000, 11200000, 400000, "Kilograms sold by guest roaster and quarter, 2026", 14, color=GREY)
    rows_a = [["Roaster", "Q1", "Q2", "Q3", "Q4"]]
    for rid, name, _ in K.ROASTERS:
        rows_a.append([name] + ["%.1f" % R["ask_a"][(name, q)] for q in (1, 2, 3, 4)])
    table(s, 457200, 1450000, 9000000, rows_a, (3000000, 1500000, 1500000, 1500000, 1500000))
    text(s, 457200, 3450000, 11000000, 900000, "Labels the counter's scales printed, counted when a card payment at the "
         "counter or a counter line at a checkout paid for them.", 13, color=GREY)
    foot(s, 5)
    # 6. funding
    s = prs.slides.add_slide(blank)
    sups = sorted(set(k[0] for k in R["ask_b"]))
    year = {sp: sum(R["ask_b"][(sp, q)] for q in (1, 2, 3, 4)) for sp in sups}
    text(s, 457200, 300000, 11200000, 700000, "%s owes the most multibuy funding for 2026" % max(year, key=year.get),
         26, True, BROWN)
    text(s, 457200, 950000, 11200000, 400000, "Funding owed by supplier and quarter of sale, EUR", 14, color=GREY)
    rows_b = [["Supplier", "Q1", "Q2", "Q3", "Q4"]]
    for sp in sups:
        rows_b.append([sp] + ["{:,}".format(int(round(R["ask_b"][(sp, q)]))) for q in (1, 2, 3, 4)])
    table(s, 457200, 1450000, 10400000, rows_b, (4400000, 1500000, 1500000, 1500000, 1500000))
    text(s, 457200, 4050000, 11000000, 900000, "Every multibuy sold in the quarter at the rate in force that day, "
         "including multibuys the tills missed and the service desk refunded later; multibuys taken back on a refund "
         "excluded.", 13, color=GREY)
    foot(s, 6)
    prs.save(path)
