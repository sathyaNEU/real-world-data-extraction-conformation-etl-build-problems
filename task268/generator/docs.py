"""The pack's documents: Faircourt's merchant service agreement (its outlets and terminals, settlement and charges),
the Corcally's Rewards terms, the coffee counter's price board and the Roastery sampling brief."""
from __future__ import annotations

import datetime as dt

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                    # noqa: E402
from reportlab.lib.pagesizes import A4                                              # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                     # noqa: E402
from reportlab.lib.units import mm                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                            # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                        # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle   # noqa: E402

import common as K   # noqa: E402

for nm, f in (("Sans", "LiberationSans-Regular"), ("Sans-B", "LiberationSans-Bold"), ("Sans-I", "LiberationSans-Italic"),
              ("Serif", "LiberationSerif-Regular"), ("Serif-B", "LiberationSerif-Bold"),
              ("Serif-I", "LiberationSerif-Italic")):
    pdfmetrics.registerFont(TTFont(nm, K.FONT_DIR + f + ".ttf"))
pdfmetrics.registerFontFamily("Sans", normal="Sans", bold="Sans-B", italic="Sans-I", boldItalic="Sans-B")
pdfmetrics.registerFontFamily("Serif", normal="Serif", bold="Serif-B", italic="Serif-I", boldItalic="Serif-B")
INK, GREY, RULE = colors.HexColor("#1d2327"), colors.HexColor("#5d666c"), colors.HexColor("#b7c0c6")
NAVY, GREEN, BROWN = colors.HexColor("#1f3a5f"), colors.HexColor("#2d5a3d"), colors.HexColor("#5a3a22")


def styles(font, accent):
    body = ParagraphStyle("b", fontName=font, fontSize=9.6, leading=13.4, textColor=INK, spaceAfter=5)
    return dict(
        body=body,
        h1=ParagraphStyle("h1", fontName=font + "-B", fontSize=15, leading=19, textColor=accent, spaceAfter=2),
        h2=ParagraphStyle("h2", fontName=font + "-B", fontSize=10.4, leading=14, textColor=INK, spaceBefore=8,
                          spaceAfter=2, keepWithNext=1),
        meta=ParagraphStyle("m", fontName=font, fontSize=8.5, leading=11.4, textColor=GREY, spaceAfter=9),
        cell=ParagraphStyle("c", fontName=font, fontSize=8.2, leading=10.2, textColor=INK),
        cellb=ParagraphStyle("cb", fontName=font + "-B", fontSize=8.2, leading=10.2, textColor=INK),
        clause=ParagraphStyle("p", parent=body, leftIndent=24, firstLineIndent=-24))


def _frame(left, right, font):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(GREY)
        canvas.setFont(font, 7.5)
        canvas.drawString(18 * mm, A4[1] - 11 * mm, left)
        canvas.drawRightString(A4[0] - 18 * mm, A4[1] - 11 * mm, right)
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _table(S, rows, widths, right=()):
    data = [[Paragraph(str(c), S["cellb"] if i == 0 else S["cell"]) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
    st = [("LINEBELOW", (0, 0), (-1, 0), .6, INK), ("LINEBELOW", (0, -1), (-1, -1), .4, RULE),
          ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]
    for c in right:
        st.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
    t.setStyle(TableStyle(st))
    return t


def _doc(path, title, author):
    return SimpleDocTemplate(path, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=19 * mm,
                             bottomMargin=17 * mm, title=title, author=author, subject="", creator="", producer="")


def _d(d):
    return "%d %s %d" % (d.day, d.strftime("%B"), d.year)


# ---------------------------------------------------------------------------------- Faircourt's merchant agreement
def write_agreement(path):
    S = styles("Serif", NAVY)
    doc = _doc(path, "Merchant Service Agreement, Schedules 2 to 4", "Faircourt Merchant Services")
    st = [Paragraph("Faircourt Merchant Services", S["h1"]),
          Paragraph("Merchant Service Agreement FMS/MSA/0447: Corcally &amp; Sons Limited. Schedules 2 to 4 as amended "
                    "to 12 May 2026. Faircourt Merchant Services Limited is authorised by the Central Bank of Ireland.",
                    S["meta"]),
          Paragraph("Schedule 2. Outlets, merchant numbers and terminals", S["h2"]),
          Paragraph("Each outlet settles under its own merchant number. A terminal is identified on every transaction "
                    "by its terminal number; a replaced terminal keeps its outlet and takes a new number from the day it "
                    "is installed.", S["body"])]
    rows = [["Outlet", "Merchant no.", "Terminal no.", "Position", "Installed", "Removed"]]
    for s, tl, tid, a, b in sorted(K.TERMINALS, key=lambda x: (x[0] != "DRM", x[0], x[1], x[3])):
        pos = "Service desk" if tl == K.DESK else "Checkout %d" % tl
        rows.append([K.STORE_NAME[s], K.MID[s], tid, pos, a.strftime("%d/%m/%Y"), b.strftime("%d/%m/%Y") if b else ""])
    rows.append(["Drumsallagh", K.MID["DRM"], K.COUNTER_TID, "Coffee counter (standalone)", "18/09/2024", ""])
    rows.append(["Online shop", K.MID["WEB"], K.WEB_TID, "Payment gateway (card not present)", "02/05/2023", ""])
    st.append(_table(S, rows, (26, 22, 22, 50, 20, 20)))
    st += [Paragraph("Schedule 3. Settlement", S["h2"]),
           Paragraph("3.1&nbsp;&nbsp;Card transactions are settled to the merchant's nominated account one banking day "
                     "after the transaction day for debit cards and two banking days after it for credit, prepaid and "
                     "commercial cards, net of the charges in Schedule 4.", S["clause"]),
           Paragraph("3.2&nbsp;&nbsp;The settlement detail file lists every transaction with its merchant and terminal "
                     "numbers, its time in Coordinated Universal Time (UTC), its type (SALE or REFUND), its amount, the "
                     "card reference (a token that stands in for the card number, so that no card number is shared with "
                     "the merchant), the card's scheme, product and region, the entry mode, the authorisation code, the "
                     "merchant's own reference where one was sent (online orders), the settlement date and the charges.",
                     S["clause"]),
           Paragraph("3.3&nbsp;&nbsp;Refunds are settled as debits on the settlement date of the refund.", S["clause"]),
           Paragraph("Schedule 4. Charges", S["h2"]),
           Paragraph("4.1&nbsp;&nbsp;Interchange is passed through at the rate the card's issuer applies and scheme fees "
                     "at cost. The service charge is EUR 0.045 per transaction plus 0.05 per cent of its amount.",
                     S["clause"]),
           Paragraph("4.2&nbsp;&nbsp;Charges are shown on each transaction in the settlement detail and summarised on the "
                     "monthly statement.", S["clause"])]
    doc.build(st, onFirstPage=_frame("Faircourt Merchant Services", "FMS/MSA/0447", "Serif"),
              onLaterPages=_frame("Faircourt Merchant Services", "FMS/MSA/0447", "Serif"))


# ---------------------------------------------------------------------------------- the Rewards terms
TERMS = [
    ("1", "The scheme", [
        ("1.1", "Corcally's Rewards is run by Corcally &amp; Sons Limited. Anyone aged 18 or over may join in store or "
                "in the Corcally's app."),
        ("1.2", "Membership is free. You may leave at any time by telling us in store or in the app.")]),
    ("2", "Your card", [
        ("2.1", "Each card is issued to one member, and a member holds one card at a time. If a card is lost or "
                "damaged we issue a new card with a new number and move the points to it."),
        ("2.2", "Cards are personal. Points earned on a card stay on that card and cannot be transferred to, or "
                "combined with, another member's card."),
        ("2.3", "Show your card or the app's barcode at the till before you pay.")]),
    ("3", "Earning points", [
        ("3.1", "You earn one point for each whole euro spent in our stores at Drumsallagh, Lisgarvan and "
                "Kilbrannagh, after any discounts."),
        ("3.2", "Points are not yet earned on online shop orders or at the coffee counter at Drumsallagh."),
        ("3.3", "No points are earned on tobacco, lottery, gift cards, mobile top-ups, prescriptions or deposit return "
                "items.")]),
    ("4", "Member prices", [
        ("4.1", "Some offers are for members. A customer without a card at the till may be given member prices at the "
                "cashier's discretion; no points are earned on that purchase.")]),
    ("5", "Using your points", [
        ("5.1", "Every 500 points becomes a EUR 5 voucher, issued at the end of each quarter."),
        ("5.2", "Vouchers can be used in our stores within three months of issue and have no cash value.")]),
    ("6", "Your information", [
        ("6.1", "We use the purchases made on your card to run the scheme, to send you offers if you have agreed to "
                "receive them, and to understand how our customers shop. Our privacy notice explains how."),
        ("6.2", "Returns are taken off your points.")]),
    ("7", "Changes", [
        ("7.1", "We may change these terms with 30 days' notice in store and in the app.")]),
]


def write_terms(path):
    S = styles("Sans", GREEN)
    doc = _doc(path, "Corcally's Rewards: terms and conditions", "Corcally & Sons Limited")
    st = [Paragraph("Corcally's Rewards", S["h1"]),
          Paragraph("Terms and conditions, January 2026 edition, in force from 5 January 2026.", S["meta"])]
    for num, head, clauses in TERMS:
        st.append(Paragraph("%s&nbsp;&nbsp;%s" % (num, head), S["h2"]))
        for c, text in clauses:
            st.append(Paragraph("%s&nbsp;&nbsp;%s" % (c, text), S["clause"]))
    doc.build(st, onFirstPage=_frame("Corcally's Rewards", "Terms, January 2026", "Sans"),
              onLaterPages=_frame("Corcally's Rewards", "Terms, January 2026", "Sans"))


# ---------------------------------------------------------------------------------- the counter's price board
def write_board(path):
    S = styles("Sans", BROWN)
    doc = _doc(path, "Coffee counter price board 2026", "Corcally's Drumsallagh")
    st = [Paragraph("Coffee counter, Drumsallagh", S["h1"]),
          Paragraph("Guest roasters, beans weighed to order and ground on request. Price board as displayed in 2026.",
                    S["meta"])]
    rows = [["Roaster", "Coffee", "Per kg, from 5 January", "Per kg, from 1 July"]]
    for rid, name, coffees in K.ROASTERS:
        for c, p1, p2 in coffees:
            rows.append([name, c, "EUR %.2f" % p1, "EUR %.2f" % p2])
    st.append(_table(S, rows, (46, 50, 38, 38), right=(2, 3)))
    st += [Spacer(1, 6), Paragraph("Accessories", S["h2"])]
    rows = [["Item", "Price"]] + [[a, "EUR %.2f" % p] for a, p in K.ACCESSORIES]
    st.append(_table(S, rows, (80, 30), right=(1,)))
    st += [Spacer(1, 6), Paragraph("Paying", S["h2"]),
           Paragraph("Card payments at the counter. Paying cash? Take your label to any checkout.", S["body"]),
           Paragraph("Open Monday to Saturday 9:00 to 18:00, Sunday 12:00 to 17:00.", S["body"])]
    doc.build(st, onFirstPage=_frame("Corcally's Drumsallagh", "Coffee counter", "Sans"),
              onLaterPages=_frame("Corcally's Drumsallagh", "Coffee counter", "Sans"))


# ---------------------------------------------------------------------------------- the sampling brief
def write_brief(path):
    from docx import Document
    from docx.shared import Pt, RGBColor
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    h = d.add_paragraph()
    r = h.add_run("Roastery sampling 2027: brief")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(18), RGBColor(0x5A, 0x3A, 0x22)
    p = d.add_paragraph("Corcally's Roastery. Alix Fennell, Roastery brand manager, 11 January 2027.")
    p.runs[0].font.size, p.runs[0].font.color.rgb = Pt(9.5), RGBColor(0x5D, 0x66, 0x6C)
    sections = [
        ("What we are doing",
         ["In spring 2027 the Roastery team will post a 227g taster pack of the House Blend, with a voucher for a "
          "full pack, to the homes that buy one rival coffee brand. The aim is to bring homes that do not have the "
          "Roastery in the cupboard to try it. This year one brand's homes get the packs."]),
        ("Which brand",
         ["The packs go to the homes of the rival whose 2026 buying homes least often also bought the Roastery: for "
          "each rival, of the homes that bought it in 2026, the share that also bought the Roastery in 2026. The lower "
          "that share, the more of the rival's homes are new to the Roastery."]),
        ("How we read 2026",
         ["A home is the people who live in it. A home bought a brand in 2026 if anyone in it bought a pack of that "
          "brand at Corcally's between 1 January and 31 December 2026, on a Corcally's Rewards card in one of our three "
          "stores or through the home's online shop account. An online order counts on the day it was delivered.",
          "The Roastery means Corcally's Roastery packs. The rivals are the other coffee brands we sell in packs. The "
          "coffee counter at Drumsallagh sells guest roasters loose and is outside this brief."]),
        ("What we need back",
         ["For each rival, the homes that bought it in 2026, how many of them also bought the Roastery, and that "
          "share; the brand the packs go to, with the runner-up and the gap between them. The deck goes to the "
          "Roastery planning session on 3 February."]),
    ]
    for head, paras in sections:
        hp = d.add_paragraph()
        hr = hp.add_run(head)
        hr.bold, hr.font.size = True, Pt(12)
        for t in paras:
            d.add_paragraph(t)
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "Corcally's Roastery"
    d.save(path)
