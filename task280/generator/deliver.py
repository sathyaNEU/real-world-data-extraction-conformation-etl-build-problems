"""The goldens: hub_peak_congestion_2026.xlsx and hub_peak_congestion_2026.png, and generator/figures.json, every
figure read back from the pack."""
from __future__ import annotations

import json
import os

import numpy as np

import analysis as A
import asks as AK
import common as K
import pack as PK


def half_up(x, nd=0):
    from decimal import ROUND_HALF_UP, Decimal
    q = Decimal(1).scaleb(-nd)
    return float(Decimal(repr(float(x))).quantize(q, rounding=ROUND_HALF_UP)) if nd else int(
        Decimal(repr(float(x))).quantize(Decimal(1), rounding=ROUND_HALF_UP))


def results(target):
    pk = PK.Pack(target)
    c = A.call(pk)
    coverage = {}
    sched = len(pk.peak_days) * (K.PEAK_TO - K.PEAK_FROM)
    for h in K.HUB:
        idx = np.where(pk.hub == h)[0]
        sel = np.isin(pk.t_hh, idx) & pk.t_peak
        coverage[h] = dict(tests=int(sel.sum()), scheduled=int(len(idx) * sched),
                           load_tested=float(pk.t_load[sel].mean()), load_all=float(pk.q_load[idx].mean()))
    return dict(pk=pk, call=c, coverage=coverage,
                credits=AK.credits(target), cap_from=AK.capacity(target, AK.CAP_FROM),
                cap_to=AK.capacity(target, AK.CAP_TO))


def figures(res):
    tab = res["call"]["table"]
    g = tab[A.GOLDEN]
    des = A.designated(g)
    hubs = {}
    for i, h in enumerate(K.HUB):
        hubs[h] = dict(name=K.HUB_NAME[h], households=int(res["call"]["n_households"][i]),
                       share=half_up(g[i]), designated=h in des,
                       outage_credit_subscriber_days=res["credits"][h],
                       capacity_14_sep_mbps=res["cap_from"][h], capacity_13_nov_mbps=res["cap_to"][h],
                       capacity_added_mbps=res["cap_to"][h] - res["cap_from"][h],
                       share_of_tests=half_up(tab["tests"][i]))
    under = [h for h in K.HUB if h not in des]
    return {
        "line": K.LINE,
        "call": {"designated": list(des), "count": len(des), "not_designated": under,
                 "nearest_over": min(des, key=lambda h: hubs[h]["share"]) if des else None,
                 "nearest_under": max(under, key=lambda h: g[K.HUB.index(h)]) if under else None},
        "hubs": hubs,
        "totals": {"households": sum(v["households"] for v in hubs.values()),
                   "outage_credit_subscriber_days": sum(v["outage_credit_subscriber_days"] for v in hubs.values()),
                   "capacity_added_mbps": sum(v["capacity_added_mbps"] for v in hubs.values())},
        "default_path": {h: half_up(tab["tests"][i]) for i, h in enumerate(K.HUB)},
        "default_path_set": list(A.designated(tab["tests"])),
        "coverage": {h: dict(load_tested=round(v["load_tested"], 4), load_all=round(v["load_all"], 4),
                             load_tested_pct=half_up(100 * v["load_tested"], 1),
                             load_all_pct=half_up(100 * v["load_all"], 1),
                             tests=v["tests"], scheduled=v["scheduled"]) for h, v in res["coverage"].items()},
        "standard_pct": int(round(100 * K.STANDARD)),
        "schedule_clauses": [1.2, 1.8, 3.1, 3.2],
    }


def chart_png(path, fig, res):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    for f in ("LiberationSans-Regular.ttf", "LiberationSans-Bold.ttf"):
        p = "/usr/share/fonts/truetype/liberation/" + f
        if os.path.exists(p):
            font_manager.fontManager.addfont(p)
    plt.rcParams.update({"font.family": "Liberation Sans", "font.size": 10})
    hubs = fig["hubs"]
    g = res["call"]["table"][A.GOLDEN]
    t = res["call"]["table"]["tests"]
    order = sorted(range(len(K.HUB)), key=lambda i: g[i])
    ink, muted, strong, pale, grid = "#1d2329", "#5a636b", "#22405e", "#b9c4cf", "#d3d9df"
    fig_, ax = plt.subplots(figsize=(8.8, 5.0), dpi=150)
    ys = list(range(len(order)))
    vals = [g[i] for i in order]
    cols = [strong if hubs[K.HUB[i]]["designated"] else pale for i in order]
    ax.barh(ys, vals, height=0.60, color=cols, edgecolor="white", linewidth=1.0, zorder=2)
    ax.scatter([t[i] for i in order], ys, marker="|", s=170, linewidths=2.0, color="#a4532c", zorder=4)
    for y, i, v in zip(ys, order, vals):
        h = K.HUB[i]
        ax.text(max(v, t[i]) + 0.5, y, "%d%%" % hubs[h]["share"], va="center", ha="left", fontsize=9.5, color=ink,
                fontweight="bold" if hubs[h]["designated"] else "normal")
    ax.set_yticks(ys)
    ax.set_yticklabels([K.HUB_NAME[K.HUB[i]] for i in order], fontsize=9.4, color=ink)
    ax.axvline(K.LINE, color=ink, linewidth=1.1, linestyle=(0, (4, 3)), zorder=3)
    ax.text(K.LINE + 0.3, len(order) - 0.45, "20% designation line (Schedule 4, 3.1)", fontsize=8.6, color=ink,
            va="bottom", ha="left")
    top = 5 * int(np.ceil((max(vals) + 4) / 5))
    ax.set_xlim(0, top)
    ax.set_ylim(-0.6, len(order) - 0.1)
    ax.set_xticks(range(0, top + 1, 5))
    ax.set_xticklabels(["%d%%" % x for x in range(0, top + 1, 5)], fontsize=8.8, color=muted)
    ax.xaxis.grid(True, color=grid, linewidth=0.6, zorder=0)
    ax.set_axisbelow(True)
    for sp in ("top", "right", "left"):
        ax.spines[sp].set_visible(False)
    ax.spines["bottom"].set_color(grid)
    ax.tick_params(axis="both", length=0)
    from matplotlib.patches import Patch
    from matplotlib.lines import Line2D
    handles = [Patch(facecolor=strong, edgecolor="none", label="Designated"),
               Patch(facecolor=pale, edgecolor="none", label="Not designated"),
               Line2D([], [], marker="|", linestyle="none", markersize=11, markeredgewidth=2.0, color="#a4532c",
                      label="Share of the completed tests below 80% of plan")]
    leg = fig_.legend(handles=handles, loc="lower left", bbox_to_anchor=(0.163, 0.105), ncol=3, fontsize=8.4,
                      frameon=False, handlelength=1.4, columnspacing=1.8)
    for txt in leg.get_texts():
        txt.set_color(muted)
    n = fig["call"]["count"]
    words = {1: "One", 2: "Two", 3: "Three", 4: "Four", 5: "Five", 6: "Six"}
    fig_.suptitle("%s of Tollard's eight hubs spent more than a fifth of the evening peak below 80%% of plan" %
                  words.get(n, str(n)), x=0.015, y=0.975, ha="left", fontsize=12.2, fontweight="bold", color=ink)
    ax.set_title("Designated under Schedule 4, 3.1 for 2027: %s" %
                 ", ".join(K.HUB_NAME[h] for h in fig["call"]["designated"]), loc="left", fontsize=9.6, color=muted,
                 pad=8)
    fig_.text(0.015, 0.02, "Source: OVCC broadband measurement panel and Tollard node utilisation filings, weekday "
                           "intervals 7 to 11 pm, 14 September to 13 November 2026.\nEach household's tests read "
                           "against its node's load in the tested interval and carried to every interval of the peak.",
              fontsize=7.4, color=muted, ha="left", linespacing=1.4)
    ax.set_xlabel("Share of the weekday evening peak below 80% of plan", fontsize=8.8, color=muted, labelpad=6)
    fig_.subplots_adjust(left=0.17, right=0.97, top=0.85, bottom=0.25)
    fig_.savefig(path, dpi=150, metadata={"Software": None})
    plt.close(fig_)


def workbook(path, fig, res):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hubs 2026"
    navy = "22405E"
    thin = Side(style="thin", color="B9C1C8")
    ws.append(["Tollard Cable hubs, 2026 measurement window: weekday evening peak below 80% of plan"])
    ws["A1"].font = Font(bold=True, size=13, color=navy)
    ws.append(["Orlenford Valley Cable Commission, performance compliance. Schedule 4, clauses 3.1 and 3.2 (designation), "
               "4.1 (capacity) and 5.1 to 5.4 (outage credits)."])
    ws["A2"].font = Font(italic=True, size=9, color="5A636B")
    ws.append([])
    head = ["Hub", "Name", "Panel households", "Share of peak below 80% of plan", "Designated for 2027",
            "Outage credits (subscriber-days)", "Capacity 14 Sep (Mbit/s)", "Capacity 13 Nov (Mbit/s)",
            "Capacity added (Mbit/s)"]
    ws.append(head)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=navy)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[4].height = 46
    g = res["call"]["table"][A.GOLDEN]
    for i, h in enumerate(K.HUB):
        d = fig["hubs"][h]
        ws.append([h, d["name"], d["households"], g[i] / 100.0, "Yes" if d["designated"] else "No",
                   d["outage_credit_subscriber_days"], d["capacity_14_sep_mbps"], d["capacity_13_nov_mbps"],
                   d["capacity_added_mbps"]])
        row = ws[ws.max_row]
        row[3].number_format = "0%"
        for c in (row[2], row[5], row[6], row[7], row[8]):
            c.number_format = "#,##0"
        if d["designated"]:
            for c in row:
                c.fill = PatternFill("solid", fgColor="E2EAF2")
            row[4].font = Font(bold=True, color=navy)
    tot = fig["totals"]
    ws.append(["All hubs", None, tot["households"], None, None, tot["outage_credit_subscriber_days"],
               sum(res["cap_from"].values()), sum(res["cap_to"].values()), tot["capacity_added_mbps"]])
    row = ws[ws.max_row]
    for c in row:
        c.font = Font(bold=True)
        c.border = Border(top=thin)
    for c in (row[2], row[5], row[6], row[7], row[8]):
        c.number_format = "#,##0"
    ws.append([])
    for line in ["Share of peak: every weekday interval from 7 to 11 pm, 14 September to 13 November 2026, of every "
                 "panel household served from the hub, each household counted equally (3.2); designated above 20% (3.1).",
                 "Outage credits: unplanned outages of four hours or more that began in the window, one day for each "
                 "started 24 hours, times the node's subscribers at the end of the month before (5.1 to 5.4).",
                 "Capacity: downstream capacity of the hub's nodes in service on the day, from the node register "
                 "filing of 16 November 2026 (4.1)."]:
        ws.append([line])
        ws.cell(ws.max_row, 1).font = Font(size=9, color="5A636B")
    for k, v in {"A": 8, "B": 15, "C": 11, "D": 13, "E": 12, "F": 15, "G": 14, "H": 14, "I": 13}.items():
        ws.column_dimensions[k].width = v
    for r in ws.iter_rows(min_row=5, max_row=13):
        for c in r[2:]:
            c.alignment = Alignment(horizontal="center" if c.column == 5 else "right")
    ws.freeze_panes = "C5"
    ws.print_area = "A1:I17"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    w2 = wb.create_sheet("Test coverage")
    w2.append(["Weekday evening peak tests against the nodes' load, 14 September to 13 November 2026"])
    w2["A1"].font = Font(bold=True, size=12, color=navy)
    w2.append([])
    w2.append(["Hub", "Tests scheduled (one an hour)", "Tests run", "Share run", "Node load in the tested intervals",
               "Node load over every peak interval", "Share of the tests below 80% of plan",
               "Share of peak below 80% of plan"])
    for c in w2[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=navy)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    w2.row_dimensions[3].height = 46
    t = res["call"]["table"]["tests"]
    for i, h in enumerate(K.HUB):
        cv = res["coverage"][h]
        w2.append([h, cv["scheduled"], cv["tests"], cv["tests"] / cv["scheduled"], cv["load_tested"], cv["load_all"],
                   t[i] / 100.0, g[i] / 100.0])
        row = w2[w2.max_row]
        row[1].number_format = "#,##0"
        row[2].number_format = "#,##0"
        for c in row[3:]:
            c.number_format = "0.0%"
    w2.append([])
    for line in ["A unit runs no test in an hour while the household's own line stays busy (methods note, 2), and the "
                 "households are busiest on the",
                 "nights the nodes carry the most traffic, so the tests that ran sit at lighter loads than the peak as "
                 "a whole. Each household's tests",
                 "were read against its node's utilisation in the tested interval (per plan, the load above which its "
                 "connection fell under 80% of plan)",
                 "and carried to every interval of the peak from the node's utilisation filing."]:
        w2.append([line])
        w2.cell(w2.max_row, 1).font = Font(size=9, color="5A636B")
    for k, v in {"A": 8, "B": 13, "C": 11, "D": 10, "E": 15, "F": 15, "G": 15, "H": 15}.items():
        w2.column_dimensions[k].width = v
    w2.freeze_panes = "B4"
    w2.page_setup.orientation = "landscape"
    w2.page_setup.fitToWidth = 1
    w2.sheet_properties.pageSetUpPr.fitToPage = True
    w2.print_area = "A1:H%d" % w2.max_row
    wb.save(path)


def write_figures(path, fig):
    with open(path, "w") as f:
        json.dump(fig, f, indent=1, sort_keys=True)
        f.write("\n")


def write(target, golden):
    os.makedirs(golden, exist_ok=True)
    res = results(target)
    fig = figures(res)
    workbook(os.path.join(golden, K.GOLDEN), fig, res)
    chart_png(os.path.join(golden, K.CHART), fig, res)
    return fig, res
