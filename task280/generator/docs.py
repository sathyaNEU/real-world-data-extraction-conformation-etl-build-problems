"""The pack's documents: the franchise agreement's Network Performance Schedule, the panel's methods note, the
Commission's data dictionary and the 2025 designation close-out (PDF); Tollard's month-end subscriber counts (XLSX);
and the pack's contents (TXT)."""
from __future__ import annotations

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                    # noqa: E402
from reportlab.lib.pagesizes import LETTER                                          # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                     # noqa: E402
from reportlab.lib.units import cm                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                            # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                        # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle   # noqa: E402

import common as K   # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("Sans", "LiberationSans-Regular"), ("Sans-B", "LiberationSans-Bold"), ("Sans-I", "LiberationSans-Italic"),
              ("Serif", "LiberationSerif-Regular"), ("Serif-B", "LiberationSerif-Bold")):
    pdfmetrics.registerFont(TTFont(nm, FONT_DIR + f + ".ttf"))
pdfmetrics.registerFontFamily("Sans", normal="Sans", bold="Sans-B", italic="Sans-I", boldItalic="Sans-B")
pdfmetrics.registerFontFamily("Serif", normal="Serif", bold="Serif-B", italic="Serif", boldItalic="Serif-B")
INK, GREY, RULE = colors.HexColor("#1d2329"), colors.HexColor("#5a636b"), colors.HexColor("#b9c1c8")
SLATE = colors.HexColor("#22405e")
PINE = colors.HexColor("#2b5743")


def styles(face="Sans", accent=SLATE):
    b = face + "-B"
    body = ParagraphStyle("b", fontName=face, fontSize=9.8, leading=13.4, textColor=INK, spaceAfter=5)
    return dict(
        body=body,
        h1=ParagraphStyle("h1", fontName=b, fontSize=14.5, leading=18, textColor=accent, spaceAfter=2),
        h2=ParagraphStyle("h2", fontName=b, fontSize=10.6, leading=14, textColor=INK, spaceBefore=8, spaceAfter=2,
                          keepWithNext=1),
        meta=ParagraphStyle("m", fontName=face, fontSize=8.4, leading=11.2, textColor=GREY, spaceAfter=9),
        cell=ParagraphStyle("c", fontName="Sans", fontSize=7.9, leading=9.9, textColor=INK),
        cellb=ParagraphStyle("cb", fontName="Sans-B", fontSize=7.9, leading=9.9, textColor=INK),
        note=ParagraphStyle("n", fontName="Sans", fontSize=7.7, leading=9.9, textColor=GREY, spaceBefore=2,
                            spaceAfter=6),
        clause=ParagraphStyle("p", parent=body, leftIndent=28, firstLineIndent=-28))


def _frame(left, right):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(GREY)
        canvas.setFont("Sans", 7.4)
        canvas.drawString(2.0 * cm, LETTER[1] - 1.25 * cm, left)
        canvas.drawRightString(LETTER[0] - 2.0 * cm, LETTER[1] - 1.25 * cm, right)
        canvas.drawRightString(LETTER[0] - 2.0 * cm, 1.1 * cm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _table(S, rows, widths, right=()):
    data = [[Paragraph(str(c), S["cellb"] if i == 0 else S["cell"]) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=[w * cm for w in widths], repeatRows=1, hAlign="LEFT")
    st = [("LINEBELOW", (0, 0), (-1, 0), .6, INK), ("LINEBELOW", (0, -1), (-1, -1), .4, RULE),
          ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 1.6),
          ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6)]
    for c in right:
        st.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
    t.setStyle(TableStyle(st))
    return t


def _clause(S, num, text):
    return Paragraph("<b>%s</b>&nbsp;&nbsp;%s" % (num, text), S["clause"])


def _doc(path, title, author, subject, left, right):
    return SimpleDocTemplate(path, pagesize=LETTER, leftMargin=2.0 * cm, rightMargin=2.0 * cm, topMargin=2.0 * cm,
                             bottomMargin=1.9 * cm, title=title, author=author, subject=subject, creator=author,
                             producer=author), _frame(left, right)


# ------------------------------------------------------------------------------------------------ the Schedule
def write_schedule(path):
    S = styles("Serif", SLATE)
    doc, fr = _doc(path, "Schedule 4: Network Performance", K.COMMISSION, "Cable franchise agreement",
                   "Cable Franchise Agreement: %s and %s" % (K.COMMISSION, K.FRANCHISEE_LEGAL),
                   "Schedule 4: Network Performance")
    st = [Paragraph("Schedule 4: Network Performance", S["h1"]),
          Paragraph("Schedule to the Cable Franchise Agreement between the %s and %s, renewed 12 June 2024 for "
                    "fifteen years. The Commission is the joint franchising authority of the eight towns of the "
                    "Orlenford Valley." % (K.COMMISSION, K.FRANCHISEE_LEGAL), S["meta"]),
          Paragraph("1. Definitions", S["h2"]),
          _clause(S, "1.1", "<b>Panel</b>: the Commission's broadband measurement panel, made up of subscriber "
                            "households that host a measurement unit on their connection."),
          _clause(S, "1.2", "<b>Panel household</b>, for a measurement window: a household of the panel whose unit was "
                            "installed before the window opened and which did not leave the panel before it closed. "
                            "A unit replaced during the window does not change the household."),
          _clause(S, "1.3", "<b>Plan</b>: the Franchisee's service plan in force for the household's account on the day. "
                            "<b>Advertised download speed</b>: the plan's download speed in the Franchisee's plan "
                            "catalogue."),
          _clause(S, "1.4", "<b>Interval</b>: a quarter-hour of local time, beginning on the hour or at 15, 30 or 45 "
                            "minutes past it."),
          _clause(S, "1.5", "<b>Weekday evening peak</b>: the intervals from 7.00 pm to 11.00 pm local time on Monday to "
                            "Friday, public holidays included."),
          _clause(S, "1.6", "<b>Measurement window</b>: from the second Monday of September to the Friday nine weeks "
                            "later, both days included."),
          _clause(S, "1.7", "<b>Delivered download speed</b>: the download throughput of the household's connection, "
                            "as the panel's download test measures it."),
          _clause(S, "1.8", "<b>Node</b>: a node of the Franchisee's network, as the node register lists it. <b>Hub</b>: "
                            "one of the eight franchise hubs of Appendix A. A household is served from the hub of the "
                            "node its service address is connected to on the day."),
          Paragraph("2. Information the Franchisee provides", S["h2"]),
          _clause(S, "2.1", "Each month the Franchisee provides the Commission with: (a) its node register; (b) the node "
                            "of each service address the Commission lists, with the dates each connection took "
                            "effect; (c) the downstream utilisation of every node serving a panel household, for each "
                            "interval from 4.00 pm to 1.00 am local time, with each interval's start stamped in "
                            "UTC; (d) its subscriber count for each node at the end of the month; and (e) its outage "
                            "reports for the month."),
          _clause(S, "2.2", "The Commission operates the panel and publishes the panel's methods."),
          Paragraph("3. Performance designation", S["h2"]),
          _clause(S, "3.1", "A hub is designated for the following year when, over the measurement window, the "
                            "connections of the panel households served from it delivered a download speed of less "
                            "than 80 per cent of the advertised download speed of their plan for more than 20 per "
                            "cent of the weekday evening peak."),
          _clause(S, "3.2", "The share in 3.1 is taken over every interval of the window's weekday evening peaks for "
                            "every panel household served from the hub, each household counted equally."),
          _clause(S, "3.3", "The Commission serves its designation notice on the Franchisee by 15 December. In each "
                            "designated hub the Franchisee adds downstream capacity before the next measurement window "
                            "opens and reports the work to the Commission by 1 September."),
          Paragraph("4. Capacity reporting", S["h2"]),
          _clause(S, "4.1", "The Franchisee reports, for each hub, the downstream capacity in service on the first and on "
                            "the last day of each measurement window. Capacity in service on a day is the sum of the "
                            "downstream capacity of the hub's nodes in service that day."),
          Paragraph("5. Outages and credits", S["h2"]),
          _clause(S, "5.1", "An outage is a loss of service on a node. It runs from the first loss of service to the "
                            "restoration of service to every subscriber served from the node."),
          _clause(S, "5.2", "For each outage of four hours or more, the Franchisee credits every subscriber served from "
                            "the node one day's charge for each started 24 hours of the outage."),
          _clause(S, "5.3", "Planned work is not an outage for 5.2: work the Franchisee notifies to the subscribers "
                            "affected at least 48 hours ahead and carries out between midnight and 6.00 am."),
          _clause(S, "5.4", "The subscribers served from a node are those in the Franchisee's subscriber count for the "
                            "node at the end of the month before the outage began."),
          Paragraph("6. Review", S["h2"]),
          _clause(S, "6.1", "The parties review this Schedule in 2029."),
          Spacer(1, 10),
          Paragraph("Appendix A. Franchise hubs", S["h2"]),
          _table(S, [["Code", "Hub", "Towns served"]] +
                 [[c, n, n] for c, n in K.HUBS], [2.0, 4.5, 9.0])]
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ------------------------------------------------------------------------------------------------ methods note
def write_methods(path, n_households):
    S = styles("Sans", PINE)
    doc, fr = _doc(path, "Panel methods note 2026", K.PEOPLE["panel"], "Broadband measurement panel", K.PANEL,
                   "Methods note 2026")
    st = [Paragraph("%s: methods note 2026" % K.PANEL, S["h1"]),
          Paragraph("%s, panel manager. Issued 21 August 2026 for the 2026 measurement window; replaces the 2025 "
                    "note." % K.PEOPLE["panel"], S["meta"]),
          Paragraph("1. The panel", S["h2"]),
          Paragraph("The panel is made up of %s subscriber households, recruited by mail and through the towns' "
                    "libraries and weighted toward no plan or hub. Each household hosts a measurement unit, "
                    "connected by cable to a LAN port of its cable modem, which the Commission's contractor installs "
                    "and maintains. Households take part for as long as they choose; a household that moves leaves "
                    "the panel." % format(n_households, ","), S["body"]),
          Paragraph("2. The download test", S["h2"]),
          Paragraph("Each unit runs one download test in every clock hour, starting at a minute drawn at random "
                    "from the first 45 minutes of the hour. The test opens eight TCP connections to one of the "
                    "Commission's test servers at the Orlenford exchange, both on the Franchisee's network, and "
                    "receives data for ten seconds. Throughput is the bytes received over the test's duration, in "
                    "kilobits per second (1 kbit/s is 1,000 bits a second).", S["body"]),
          Paragraph("Before a test the unit measures the household's own traffic on the line. If more than 64 "
                    "kbit/s is flowing downstream, it waits five minutes and measures again, up to three more "
                    "times; if the line is still in use at the last check, it runs no test in that hour. A unit "
                    "records only the tests it runs.", S["body"]),
          Paragraph("3. Time", S["h2"]),
          Paragraph("Units keep local time from the network and stamp each test with its local start time and the "
                    "offset from UTC in force at that moment (UTC-04:00 under daylight time, UTC-05:00 from 2.00 am "
                    "on 1 November 2026).", S["body"]),
          Paragraph("4. Delivery", S["h2"]),
          Paragraph("Units upload their results every hour. A unit that loses its link to the collector keeps its "
                    "results and uploads them when the link returns; the collector can then receive the same "
                    "results twice. A failed unit is replaced, under a new unit number, at the same household.",
                    S["body"]),
          Paragraph("5. Latency", S["h2"]),
          Paragraph("Units also send a small UDP probe to the test servers every second and keep a daily summary of "
                    "the round-trip times and losses between 5.00 pm and midnight.", S["body"]),
          Paragraph("6. The evening extract", S["h2"]),
          Paragraph("For each measurement window the Commission draws an extract of every download test the units "
                    "ran between 5.00 pm and midnight local time on every day of the window, with the panel's "
                    "household and unit records as they stand at the extract.", S["body"]),
          Paragraph("Contact: %s, panel manager." % K.PEOPLE["panel"], S["body"])]
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ------------------------------------------------------------------------------------------------ data dictionary
DICTIONARY = [
    ("panel_download_tests_2026_evening.csv", "One row a download test a panel unit ran between 5.00 pm and midnight "
                                              "local time, 14 September to 13 November 2026.", [
        ("unit_id", "the measurement unit"),
        ("test_start_local, utc_offset", "the unit's local start time and the offset from UTC in force"),
        ("target_server", "the test server"),
        ("threads, duration_ms", "TCP connections opened and the test's duration in milliseconds"),
        ("bytes_received, throughput_kbps", "bytes received and throughput in kbit/s")]),
    ("panel_units.csv", "One row a measurement unit the panel has installed.", [
        ("household_id", "the household the unit is installed at"),
        ("installed_on, removed_on", "dates the unit went in and came out"),
        ("removal_reason", "why it came out")]),
    ("panel_households.csv", "One row a household that has taken part in the panel.", [
        ("tollard_account", "the household's account number with the Franchisee"),
        ("service_address", "the Franchisee's service address number, as the household gave it at enrolment"),
        ("enrolled_on, withdrawn_on", "dates the household joined and left the panel"),
        ("plan_at_enrolment", "the plan the household named at enrolment")]),
    ("tollard_plan_catalogue.csv", "One row a plan in the Franchisee's catalogue.", [
        ("download_mbps, upload_mbps", "advertised speeds, Mbit/s"),
        ("launched_on, withdrawn_on", "dates the plan was offered from and to"),
        ("available_on", "the nodes the plan is sold on")]),
    ("tollard_subscription_history.csv", "One row a plan held on a panel household's account, from the Franchisee's "
                                         "billing system.", [
        ("effective_from, effective_to", "first and last day the plan was in force; blank to while current")]),
    ("tollard_node_register.csv", "One row a node record in the Franchisee's node register (filing of 16 November "
                                  "2026).", [
        ("hub_code", "the franchise hub"),
        ("node_class, ds_capacity_mbps", "the node's equipment and its downstream capacity, Mbit/s"),
        ("in_service_from, in_service_to", "first and last day of service in this record; blank to while in service"),
        ("split_from", "the node this node was split from")]),
    ("tollard_address_node_table.csv", "One row a connection of a listed service address to a node, from the "
                                       "Franchisee's filing.", [
        ("effective_from, effective_to", "first and last day of the connection; blank to while current")]),
    ("tollard_node_utilisation_2026q3q4.csv", "One row a node serving a panel household and an interval from 4.00 pm "
                                             "to 1.00 am local time, 14 September to 13 November 2026.", [
        ("interval_start_utc", "start of the interval, UTC"),
        ("ds_util_pct", "downstream traffic as a share of the node's downstream capacity over the interval, per cent")]),
    ("tollard_node_subscriber_counts_2026.xlsx", "The Franchisee's subscriber count for each node at each month "
                                                 "end, July to October 2026.", []),
    ("tollard_outage_reports_2026.csv", "One row a node named in one of the Franchisee's outage reports for August "
                                        "to November 2026.", [
        ("ticket_id", "the report's ticket"),
        ("outage_start_utc, service_restored_utc", "UTC"),
        ("work_type", "Unplanned, or Planned work notified under Schedule 4, 5.3"),
        ("cause_code", "the Franchisee's cause code")]),
    ("ovcc_consumer_complaints_2026.csv", "One row a complaint about the Franchisee the Commission received, July "
                                          "to November 2026.", []),
    ("panel_latency_2026_evening.csv", "One row a unit's daily latency summary, 5.00 pm to midnight local time.", [
        ("rtt_median_ms, rtt_p95_ms", "median and 95th percentile round-trip time"),
        ("loss_pct", "probes lost, per cent")]),
]


def write_dictionary(path):
    S = styles("Sans", SLATE)
    doc, fr = _doc(path, "Panel data dictionary", K.PEOPLE["panel"], "Broadband measurement panel", K.COMMISSION,
                   "Data dictionary")
    st = [Paragraph("%s: data dictionary" % K.COMMISSION, S["h1"]),
          Paragraph("%s, panel manager. Revised 18 November 2026." % K.PEOPLE["panel"], S["meta"])]
    for name, grain, cols in DICTIONARY:
        st.append(Paragraph(name, S["h2"]))
        st.append(Paragraph(grain, S["body"]))
        if cols:
            st.append(_table(S, [["Column", "Meaning"]] + [list(c) for c in cols], [5.4, 11.6]))
            st.append(Spacer(1, 4))
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ------------------------------------------------------------------------------------------------ 2025 close-out
def write_closeout(path, hubs, works):
    """hubs: the 2025 designations; works: (hub, node, work, completed, capacity added Mbit/s)"""
    S = styles("Sans", SLATE)
    doc, fr = _doc(path, "Designation close-out 2025", K.PEOPLE["director"], "Network performance", K.COMMISSION,
                   "Designation close-out: 2025 window")
    names = " and ".join(K.HUB_NAME[h] for h in hubs)
    rows = [["Hub", "Node", "Work", "Completed", "Capacity added (Mbit/s)"]]
    tot = {}
    for h, n, w, done, cap in works:
        rows.append([K.HUB_NAME[h], n, w, done.strftime("%d %b %Y").lstrip("0"), format(cap, ",")])
        tot[h] = tot.get(h, 0) + cap
    st = [Paragraph("Designation close-out: 2025 measurement window", S["h1"]),
          Paragraph("%s, executive director. 11 September 2026." % K.PEOPLE["director"], S["meta"]),
          Paragraph("1. Designations", S["h2"]),
          Paragraph("On 12 December 2025 the Commission designated %s under Schedule 4, 3.1, on the panel's 2025 "
                    "measurement window. The Franchisee's remedy plan, agreed on 20 February 2026, split or "
                    "upgraded the nodes below." % names, S["body"]),
          Paragraph("2. Works", S["h2"]),
          _table(S, rows, [3.0, 2.2, 5.6, 2.8, 3.4], right=(4,)),
          Paragraph("Source: the Franchisee's node register (filing of 16 August 2026) and its completion report of "
                    "28 August 2026.", S["note"]),
          Paragraph("3. Close-out", S["h2"]),
          Paragraph("The Commission's engineer inspected the works at the hub sites on 3 and 4 September 2026. The "
                    "remedy for %s is closed; the hubs are measured again in the 2026 window like every other hub. "
                    "Capacity added: %s." % (names, "; ".join("%s %s Mbit/s" % (K.HUB_NAME[h], format(v, ","))
                                                              for h, v in tot.items())), S["body"])]
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ------------------------------------------------------------------------------------------------ subscriber counts
def write_subcounts(path, head, rows):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Subscribers by node"
    ws.append(head)
    for r in rows:
        ws.append(list(r))
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="22405E")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "C2"
    ws.column_dimensions["A"].width = 11
    ws.column_dimensions["B"].width = 9
    for col in "CDEF":
        ws.column_dimensions[col].width = 21
    for row in ws.iter_rows(min_row=2, min_col=3):
        for c in row:
            c.number_format = "#,##0"
    w2 = wb.create_sheet("Notes")
    for line in ["%s: subscriber count by node at each month end, filed under Schedule 4, 2.1(d)." % K.FRANCHISEE_LEGAL,
                 "A subscriber is an account with active service on the node at 11.59 pm local time on the month's "
                 "last day. A node not in service at a month end has no count for it.",
                 "Prepared by %s, director of regulatory affairs, 9 November 2026." % K.PEOPLE["tollard"]]:
        w2.append([line])
    w2.column_dimensions["A"].width = 110
    wb.save(path)


def write_contents(path, items):
    out = ["%s: 2026 network performance review pack" % K.COMMISSION,
           "Assembled %s by the Commission's performance team." % K.ASSEMBLED.strftime("%d %B %Y").lstrip("0"), ""]
    for name, desc in items:
        out.append(name)
        out.append("    " + desc)
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(out) + "\n")
