"""The pack's documents: the Relief Coverage Standard, the Director's allowance memo and the protective services
agreement excerpts (PDF); and the Comptroller's overtime expenditure table (XLSX)."""
from __future__ import annotations

import os

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                    # noqa: E402
from reportlab.lib.pagesizes import LETTER                                          # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                     # noqa: E402
from reportlab.lib.units import inch                                                # noqa: E402
from reportlab.pdfbase import pdfmetrics                                            # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                        # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle   # noqa: E402

import common as K   # noqa: E402

FONT_DIR = next((d for d in ("/usr/share/fonts/truetype/liberation/", os.path.expanduser("~/Library/Fonts/"),
                              "/Library/Fonts/") if os.path.exists(d + "LiberationSerif-Regular.ttf")),
                "/usr/share/fonts/truetype/liberation/")
for nm, f in (("Sans", "LiberationSans-Regular"), ("Sans-B", "LiberationSans-Bold"), ("Sans-I", "LiberationSans-Italic"),
              ("Serif", "LiberationSerif-Regular"), ("Serif-B", "LiberationSerif-Bold"),
              ("Serif-I", "LiberationSerif-Italic")):
    pdfmetrics.registerFont(TTFont(nm, FONT_DIR + f + ".ttf"))
pdfmetrics.registerFontFamily("Sans", normal="Sans", bold="Sans-B", italic="Sans-I", boldItalic="Sans-B")
pdfmetrics.registerFontFamily("Serif", normal="Serif", bold="Serif-B", italic="Serif-I", boldItalic="Serif-B")
INK, GREY, RULE = colors.HexColor("#1b2430"), colors.HexColor("#56616c"), colors.HexColor("#b8c0c8")
NAVY = colors.HexColor("#1f3b57")
MAROON = colors.HexColor("#5a2330")


def styles(face="Sans", accent=NAVY):
    b = face + "-B"
    body = ParagraphStyle("b", fontName=face, fontSize=10.2, leading=13.9, textColor=INK, spaceAfter=5)
    return dict(
        body=body,
        h1=ParagraphStyle("h1", fontName=b, fontSize=14.5, leading=18, textColor=accent, spaceAfter=2),
        h2=ParagraphStyle("h2", fontName=b, fontSize=10.8, leading=14, textColor=INK, spaceBefore=8, spaceAfter=2,
                          keepWithNext=1),
        meta=ParagraphStyle("m", fontName=face, fontSize=8.6, leading=11.4, textColor=GREY, spaceAfter=9),
        cell=ParagraphStyle("c", fontName="Sans", fontSize=8.2, leading=10.2, textColor=INK),
        cellb=ParagraphStyle("cb", fontName="Sans-B", fontSize=8.2, leading=10.2, textColor=INK),
        note=ParagraphStyle("n", fontName="Sans", fontSize=7.9, leading=10.1, textColor=GREY, spaceBefore=2,
                            spaceAfter=6),
        memo=ParagraphStyle("mh", fontName="Sans", fontSize=10.2, leading=14.6, textColor=INK),
        clause=ParagraphStyle("p", parent=body, leftIndent=30, firstLineIndent=-30),
        sub=ParagraphStyle("s", parent=body, leftIndent=48, firstLineIndent=-18))


def _frame(left, right):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(GREY)
        canvas.setFont("Sans", 7.6)
        canvas.drawString(0.9 * inch, LETTER[1] - 0.55 * inch, left)
        canvas.drawRightString(LETTER[0] - 0.9 * inch, LETTER[1] - 0.55 * inch, right)
        canvas.drawRightString(LETTER[0] - 0.9 * inch, 0.5 * inch, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _table(S, rows, widths, right=()):
    data = [[Paragraph(str(c), S["cellb"] if i == 0 else S["cell"]) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=[w * inch for w in widths], repeatRows=1, hAlign="LEFT")
    st = [("LINEBELOW", (0, 0), (-1, 0), .6, INK), ("LINEBELOW", (0, -1), (-1, -1), .4, RULE),
          ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 1.8),
          ("BOTTOMPADDING", (0, 0), (-1, -1), 1.8)]
    for c in right:
        st.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
    t.setStyle(TableStyle(st))
    return t


def _clause(S, num, text, key="clause"):
    return Paragraph("<b>%s</b>&nbsp;&nbsp;%s" % (num, text), S[key])


def _doc(path, title, author, subject, left, right):
    return SimpleDocTemplate(path, pagesize=LETTER, leftMargin=0.9 * inch, rightMargin=0.9 * inch,
                             topMargin=0.85 * inch, bottomMargin=0.8 * inch, title=title, author=author,
                             subject=subject, creator=author, producer=author), _frame(left, right)


# ------------------------------------------------------------------------------------------------ the standard
def write_standard(path):
    S = styles("Serif", NAVY)
    doc, fr = _doc(path, "Relief Coverage Standard for Round-the-Clock Units", "Office of Budget Analysis",
                   "Budget Instruction 2024-07", "%s, %s" % (K.STATE, K.DBM), "Budget Instruction 2024-07")
    st = [Paragraph("Relief Coverage Standard for Round-the-Clock Units", S["h1"]),
          Paragraph("Budget Instruction 2024-07. Office of Budget Analysis, %s. Issued 14 August 2024; applies "
                    "to the allowances from FY2026. It replaces the annual relief factor of Budget Instruction "
                    "2016-03, which sized complements on a year's average hours." % K.DBM, S["meta"]),
          Paragraph("1. Purpose", S["h2"]),
          _clause(S, "1.1", "This standard sizes the relief complement the allowance funds for a unit that staffs "
                            "fixed posts around the clock on lines. A line is one officer's rota of five eight-hour "
                            "shifts a week on one of the unit's shifts; the officer who holds it is its holder."),
          _clause(S, "1.2", "Relief positions fill the posts whose lines their holders do not work, so that those "
                            "posts are not filled on overtime or left short."),
          Paragraph("2. Definitions", S["h2"]),
          _clause(S, "2.1", "The year planned for is the fiscal year of the allowance in preparation. The last "
                            "closed year is the latest fiscal year whose rosters the timekeeping systems have "
                            "closed."),
          _clause(S, "2.2", "A week is a fiscal week: seven days counted from 1 July. The day or two after a "
                            "year's 52nd week belong to no week."),
          _clause(S, "2.3", "A unit's backfill need in a week is the scheduled hours on its lines in that week "
                            "that their holders will not work on post:"),
          _clause(S, "(a)", "every scheduled hour a holder is off post, whether on leave of any kind, injured on "
                            "duty or in training away from the post; and", "sub"),
          _clause(S, "(b)", "every scheduled hour of a line the position control carries without a holder, until "
                            "the first shift on post of the officer who fills it.", "sub"),
          _clause(S, "2.4", "A unit's lines and their holders are those of the position control for the year "
                            "planned for. A line that falls vacant during the year is backfilled from the unit's "
                            "salary savings while the agency recruits to it; the relief complement is not sized "
                            "for it."),
          _clause(S, "2.5", "A relief position is established by title, for one unit, and carries no line. Its "
                            "holder is scheduled for 40 hours a week to whichever of the unit's posts need filling."),
          Paragraph("3. Sizing a unit's relief complement", S["h2"]),
          _clause(S, "3.1", "A unit's relief complement meets the unit's backfill need in full in every week of "
                            "the year planned for but four."),
          _clause(S, "3.2", "The complement is the smallest whole number of relief positions whose hours on post "
                            "in a week reach the unit's backfill need in its fifth-heaviest week of the year "
                            "planned for."),
          _clause(S, "3.3", "A relief position's hours on post in a week are the hours its holder will work on "
                            "post in an average week of the year planned for: the 40 scheduled hours, less the "
                            "hours the holder will be off post in the sense of 2.3(a)."),
          _clause(S, "3.4", "The off-post hours of a unit's lines in each week of the year planned for are those "
                            "its lines carried in the same fiscal week of the last closed year."),
          Paragraph("4. Workings", S["h2"]),
          _clause(S, "4.1", "The Office keeps each unit's workings with the allowance file: the weekly backfill "
                            "need, the hours a relief position works on post, and the complement."),
          Spacer(1, 8),
          Paragraph("Questions on this instruction go to the Office of Budget Analysis, workforce analysis desk.",
                    S["note"])]
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ------------------------------------------------------------------------------------------------ the memo
def write_memo(path):
    S = styles("Sans", MAROON)
    doc, fr = _doc(path, "FY2028 allowance: overtime-relief positions", K.PEOPLE["director"],
                   "FY2028 allowance", "%s, %s" % (K.STATE, K.DBM), "Office of Budget Analysis")
    head = _table(S, [["", ""],
                      ["To", "%s, Lead, Workforce Analysis Desk, Office of Budget Analysis" % K.PEOPLE["desk"]],
                      ["From", "%s, Director, Office of Budget Analysis" % K.PEOPLE["director"]],
                      ["Date", "9 October 2026"],
                      ["Subject", "FY2028 allowance: overtime-relief positions for the round-the-clock units"]],
                  [0.9, 5.6])
    rows = [["Unit", "Agency", "Title"]] + [[K.UNIT_NAME[u], K.AGENCY[K.UNIT_TITLE[u]], K.TITLES[K.UNIT_TITLE[u]]]
                                            for u in K.UNIT_CODES]
    st = [Paragraph("Memorandum", S["h1"]), head, Spacer(1, 10),
          Paragraph("The Governor's FY2028 allowance will fund 60 overtime-relief positions for the eight "
                    "round-the-clock units below, as new positions from 1 July 2027. Please place them under the "
                    "Relief Coverage Standard (Budget Instruction 2024-07), on the FY2026 rosters, our last closed "
                    "year, and the FY2028 position control, and send me each unit's complement by 6 November 2026; "
                    "I will enter the placement in the agencies' allowances.", S["body"]),
          Paragraph("The positions will be filled from the agencies' July 2027 entry classes. %s, director of "
                    "recruitment and training at Public Safety, will size those classes once the placement is "
                    "settled, and Health and Transportation have agreed to work from the same placement."
                    % K.PEOPLE["recruit"], S["body"]),
          Paragraph("If the units need fewer than 60, the rest return to the Governor's reserve. If they need "
                    "more, tell me how many more and I will take it to the Secretary before the allowance closes; "
                    "the 60 are not to be spread thinner than the standard allows.", S["body"]),
          Paragraph("For the briefing that goes with the allowance I will also need the overtime worked at these "
                    "units in FY2026 and the people who worked it. %s at Central Payroll has sent the overtime "
                    "lines from the checks of July 2025 to July 2026." % K.PEOPLE["payroll"], S["body"]),
          Spacer(1, 4),
          _table(S, rows, [2.7, 2.9, 1.6]),
          Spacer(1, 10),
          Paragraph("LW/ca", S["note"])]
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ------------------------------------------------------------------------------------------------ the agreement
def write_agreement(path):
    S = styles("Serif", NAVY)
    doc, fr = _doc(path, "Memorandum of Understanding 2025-2028: excerpts", "Office of Personnel Services",
                   "Protective Services bargaining unit", "%s and the Merrimont Protective Services Association"
                   % K.STATE, "MOU 2025-2028, excerpts")
    st = [Paragraph("Memorandum of Understanding, 1 July 2025 to 30 June 2028", S["h1"]),
          Paragraph("Between the %s and the Merrimont Protective Services Association (Correctional Officers, "
                    "Police Communications Operators and Highway Maintenance Workers). "
                    "Excerpts: Articles 14, 15, 16 and 32, and the side letter of 27 March 2026." % K.STATE,
                    S["meta"]),
          Paragraph("Article 14. Overtime", S["h2"]),
          _clause(S, "14.1", "Hours worked beyond 40 in a work week, or on a scheduled day off, are overtime and "
                             "are paid at one and one-half times the employee's base hourly rate."),
          _clause(S, "14.2", "An employee may elect compensatory time instead of overtime pay. Compensatory time "
                             "is credited at one and one-half hours for each hour of overtime worked, to a balance "
                             "of no more than 480 hours."),
          _clause(S, "14.3", "Overtime at a unit is offered by seniority from the unit's voluntary list, which "
                             "employees of the same title at the State's other units may also join. When the list "
                             "is exhausted, the unit holds over its junior employee on duty or calls in its junior "
                             "employee off duty, for a full shift or for four hours."),
          Paragraph("Article 15. Shift differential", S["h2"]),
          _clause(S, "15.1", "An employee who works between 11:00 p.m. and 7:00 a.m. receives a night "
                             "differential of $1.25 for each hour so worked, overtime hours included."),
          Paragraph("Article 16. Hours of overtime", S["h2"]),
          _clause(S, "16.1", "No employee shall work more than 600 hours of overtime in a fiscal year without the "
                             "written approval of the appointing authority, which the employer reports to the "
                             "Association each quarter."),
          Paragraph("Article 32. Wages", S["h2"]),
          _clause(S, "32.1", "Base hourly rates are those of the State's Protective Services salary schedule, by "
                             "grade and step."),
          _clause(S, "32.2", "Steps advance on 1 July for employees with a satisfactory rating."),
          Paragraph("Side letter of 27 March 2026", S["h2"]),
          Paragraph("The parties agree a general increase of 3.0 per cent to the Protective Services salary "
                    "schedule, effective with the pay period beginning 29 June 2025. Central Payroll will load "
                    "the new rates for the pay period beginning 5 April 2026 and will pay the increase on "
                    "earnings of the earlier pay periods as retroactive pay on the check of 24 April 2026.",
                    S["body"]),
          Paragraph("For the State: %s, chief negotiator, Office of Personnel Services. For the Association: %s, "
                    "president." % (K.PEOPLE["negotiator"], K.PEOPLE["union"]), S["note"])]
    doc.build(st, onFirstPage=fr, onLaterPages=fr)


# ------------------------------------------------------------------------------------------------ OT by program
def write_otprog(path, rows):
    """rows: (agency, program, overtime hours, overtime pay)"""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Table 7"
    ws.append(["Table 7. Overtime expenditure by agency and program, FY2026"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Checks dated 1 July 2025 to 30 June 2026. All funds."])
    ws["A2"].font = Font(italic=True, size=9, color="56616C")
    ws.append([])
    ws.append(["Agency", "Program", "Overtime hours paid", "Overtime pay ($)"])
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3B57")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(list(r))
        ws.cell(ws.max_row, 3).number_format = "#,##0"
        ws.cell(ws.max_row, 4).number_format = "#,##0"
    ws.append([])
    ws.append(["Source: Comptroller of Merrimont, FY2026 Annual Personnel Expenditure Report."])
    ws.cell(ws.max_row, 1).font = Font(size=8.5, color="56616C")
    for k, v in {"A": 48, "B": 44, "C": 16, "D": 16}.items():
        ws.column_dimensions[k].width = v
    ws.freeze_panes = "A5"
    wb.save(path)
