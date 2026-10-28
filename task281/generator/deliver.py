"""The goldens for task281, computed from the shipped pack by analysis.py: the 2027 program workbook, the projection
chart and generator/figures.json (every figure the submission and the checks quote)."""
import json
import os

import numpy as np

import analysis as A
import params as P

AT_RISK = P.DRV["A"]
CURED = P.DRV["B"]


def nearest(v, u):
    return int(round(v / u) * u)


def compute(target):
    D = A.load(target)
    G = A.golden_call(D)
    rows = G["rows"]
    E, _ = A.energy_saving(D)
    late, tot, _ = A.repairs(D)
    pooled, (kP, lP) = A.rung_pooled(rows)
    expB = float(sum((r["t"] / lP) ** kP for r in rows if r["maker"] == CURED))
    ph = {}
    for pc in P.PHASE_CODES:
        v = G["phases"][pc]
        ph[pc] = dict(installed=v["installed"], failed_to_date=v["failed_to_date"], on_original=v["n"],
                      at_risk_maker=v["by_maker"][AT_RISK], cured_maker=v["by_maker"][CURED],
                      share=round(v["share"], 6), share5=nearest(v["share"] * 100, 5), designated=v["designated"],
                      share_pct=round(v["share"] * 100, 1), pooled_share=round(pooled[pc]["share"], 6),
                      pooled_pct=round(pooled[pc]["share"] * 100, 1))
    program = sum(ph[pc]["on_original"] for pc in G["set"])
    fits = {mk: dict(drivers=v["n"], failures=v["failures"], shape=round(v["k"], 4), shape1=round(v["k"], 1),
                     life1=round(float(v["lam"]), 1),
                     life=round(float(v["lam"]), 4), F12=round(v["F12"], 6)) for mk, v in G["fits"].items()}
    names = ", ".join(G["set"][:-1]) + " and " + G["set"][-1]
    fig = dict(
        constants=dict(trigger_pct=45, repair_standard_hours=72, reprogrammed_watts=70, horizon_years=12, share_unit_pct=5,
                       clauses_cited=["2.1", "2.2", "2.4", "3.1", "3.2", "4.2"]),
        call=dict(set=G["set"], program_luminaires=program,
                  default_luminaires=sum(ph[pc]["on_original"] for pc in P.PHASE_CODES),
                  sentence="Put conversion phases %s into the 2027 group driver replacement program, %s luminaires "
                           "re-driven." % (names, format(program, ","))),
        phases=ph, fits=fits,
        pooled=dict(shape=round(kP, 4), life=round(lP, 4), expected_cured_failures=round(expB, 1),
                    expected_floor=int(expB // 100 * 100),
                    set=[pc for pc in P.PHASE_CODES if pooled[pc]["share"] > P.TRIGGER]),
        towns={t: dict(energy_mwh=round(E[t], 3), energy=int(round(E[t])), repaired=tot[t], late=late[t])
               for t in sorted(E)},
    )
    return D, G, fig


# ---------------------------------------------------------------------------------------------- the workbook
def write_xlsx(path, fig):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "2027 program"
    bold = Font(name="Calibri", bold=True)
    head_fill = PatternFill("solid", fgColor="DCE6EA")
    thin = Side(style="thin", color="9AA6B0")
    names = ", ".join(fig["call"]["set"][:-1]) + " and " + fig["call"]["set"][-1]
    ws["A1"] = "2027 group driver replacement: %s go in, %s luminaires" % (names,
                                                                           format(fig["call"]["program_luminaires"], ","))
    ws["A1"].font = Font(name="Calibri", bold=True, size=13)
    ws["A2"] = ("%s street lighting. Program test of Program Note SL-LM-04 rev. 3: work-management record closed "
                "31 December 2026, luminaires on their original driver on 1 January 2027, failure before the "
                "twelfth anniversary of installation." % P.UTILITY)
    ws["A2"].font = Font(name="Calibri", italic=True, size=9, color="56616C")
    head = ["Phase", "Installed", "Original-driver failures to 31 Dec 2026", "On original driver 1 Jan 2027",
            "of which %s" % AT_RISK, "of which %s" % CURED, "Projected share failing before year 12",
            "2027 program"]
    ws.append([])
    ws.append(head)
    for c in ws[4]:
        c.font = bold
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(bottom=thin)
    for pc in P.PHASE_CODES:
        v = fig["phases"][pc]
        ws.append([pc, v["installed"], v["failed_to_date"], v["on_original"], v["at_risk_maker"], v["cured_maker"],
                   v["share5"] / 100.0, "Yes" if v["designated"] else "No"])
        r = ws.max_row
        for col in "BCDEF":
            ws["%s%d" % (col, r)].number_format = "#,##0"
        ws["G%d" % r].number_format = "0%"
        ws["H%d" % r].alignment = Alignment(horizontal="center")
    ws.append([])
    ws.append(["Program", None, None, fig["call"]["program_luminaires"], None, None, None,
               "%d phases" % len(fig["call"]["set"])])
    r = ws.max_row
    ws["A%d" % r].font = bold
    ws["D%d" % r].number_format = "#,##0"
    ws["D%d" % r].font = bold
    ws.append([])
    ws.append(["Shares to the nearest 5 per cent; a phase enters above 45 per cent (note 2.1). Each luminaire is "
               "projected under its own driver maker's lifetime (sheet Driver lifetimes), the maker read from its "
               "delivery lot's certificate of conformance."])
    ws["A%d" % ws.max_row].font = Font(name="Calibri", italic=True, size=9, color="56616C")
    ws.append(["Source: Brayle Electric installation logs 2021 to 2025; Halberd certificates of conformance; legacy "
               "WMS export to June 2023 and EAM work orders to 31 December 2026."])
    ws["A%d" % ws.max_row].font = Font(name="Calibri", italic=True, size=9, color="56616C")
    ws.print_title_rows = "4:4"
    for col, wdt in zip("ABCDEFGH", [8, 10, 15, 15, 14, 14, 15, 11]):
        ws.column_dimensions[col].width = wdt
    ws.row_dimensions[4].height = 46
    ws.freeze_panes = "B5"

    w2 = wb.create_sheet("Driver lifetimes")
    w2["A1"] = "Original-driver lifetimes by driver maker, Weibull by maximum likelihood with suspensions"
    w2["A1"].font = Font(name="Calibri", bold=True, size=12)
    w2.append([])
    w2.append(["Driver maker", "Drivers installed", "Failures to 31 Dec 2026", "Weibull shape",
               "Characteristic life (years)", "Failed by year 12"])
    for c in w2[3]:
        c.font = bold
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    v = fig["fits"][AT_RISK]
    w2.append([AT_RISK, v["drivers"], v["failures"], v["shape1"], v["life1"], v["F12"]])
    r = w2.max_row
    w2["B%d" % r].number_format = "#,##0"
    w2["D%d" % r].number_format = "0.0"
    w2["E%d" % r].number_format = "0.0"
    w2["F%d" % r].number_format = "0%"
    v = fig["fits"][CURED]
    w2.append([CURED, v["drivers"], v["failures"], "no wear-out", None, "under 1%"])
    w2["B%d" % w2.max_row].number_format = "#,##0"
    w2["F%d" % w2.max_row].alignment = Alignment(horizontal="right")
    w2.append([])
    pl = fig["pooled"]
    w2.append(["One Weibull over every conversion driver expects more than %s failures among %s' drivers by "
               "31 December 2026, against the %d recorded." % (format(pl["expected_floor"], ","), CURED,
                                                               fig["fits"][CURED]["failures"])])
    w2["A%d" % w2.max_row].alignment = Alignment(wrap_text=False)
    for col, wdt in zip("ABCDEF", [26, 13, 13, 11, 14, 11]):
        w2.column_dimensions[col].width = wdt
    w2.row_dimensions[3].height = 32

    w3 = wb.create_sheet("Towns 2026")
    w3["A1"] = "Converted poles' 2026 energy saving and the 72-hour repair standard, by town"
    w3["A1"].font = Font(name="Calibri", bold=True, size=12)
    w3.append([])
    w3.append(["Town", "Energy saved against the sodium fixtures (MWh)", "Outages repaired in 2026",
               "Missed the 72-hour standard"])
    for c in w3[3]:
        c.font = bold
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for t, v in fig["towns"].items():
        w3.append([t, v["energy"], v["repaired"], v["late"]])
        w3["B%d" % w3.max_row].number_format = "#,##0"
    w3.append([])
    w3.append(["Energy on Schedule SL's accounting (3.1, 3.2), each converted pole against the sodium fixture the "
               "installer took down, whole MWh. Repairs timed under 4.2 from the first report of the outage, any "
               "light on an underground circuit counting for a cable fault, or from the Company's own find."])
    w3["A%d" % w3.max_row].font = Font(name="Calibri", italic=True, size=9, color="56616C")
    w3.append(["Source: Schedule SL; Corrinvale networked lighting control records 2026; customer outage reports, "
               "December 2025 to December 2026; GIS register at 31 December 2026."])
    w3["A%d" % w3.max_row].font = Font(name="Calibri", italic=True, size=9, color="56616C")
    for col, wdt in zip("ABCD", [18, 20, 14, 14]):
        w3.column_dimensions[col].width = wdt
    w3.row_dimensions[3].height = 46
    wb.save(path)


# ---------------------------------------------------------------------------------------------- the chart
def write_png(path, fig):
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib import font_manager
    import matplotlib.pyplot as plt
    fd = next((d for d in ("/usr/share/fonts/truetype/liberation/", os.path.expanduser("~/Library/Fonts/"),
                           "/Library/Fonts/") if os.path.exists(d + "LiberationSans-Regular.ttf")), None)
    if fd:
        for f in ("LiberationSans-Regular.ttf", "LiberationSans-Bold.ttf"):
            font_manager.fontManager.addfont(fd + f)
        plt.rcParams["font.family"] = "Liberation Sans"
    plt.rcParams["svg.hashsalt"] = "task281"
    ph = fig["phases"]
    x = np.arange(len(P.PHASE_CODES))
    vals = [ph[pc]["share"] * 100 for pc in P.PHASE_CODES]
    cols = ["#1d4e5a" if ph[pc]["designated"] else "#a9b6bf" for pc in P.PHASE_CODES]
    f_, ax = plt.subplots(figsize=(8.6, 4.6), dpi=110)
    ax.bar(x, vals, color=cols, width=0.62, zorder=2)
    ax.axhline(45, color="#9b3a1f", lw=1.4, ls="--", zorder=3)
    ax.text(len(x) - 0.45, 46.2, "Program trigger, 45%", color="#9b3a1f", ha="right", va="bottom", fontsize=9)
    for i, pc in enumerate(P.PHASE_CODES):
        ax.text(i, vals[i] + 1.0, "%d%%" % ph[pc]["share5"], ha="center", va="bottom", fontsize=9,
                color="#1b2430")
    ax.set_xticks(x)
    ax.set_xticklabels(P.PHASE_CODES)
    ax.set_ylim(0, 80)
    ax.set_ylabel("Projected share failing before year 12 (%)")
    ax.set_xlabel("Conversion phase (installed January 2021 to September 2025)")
    ax.set_title("Five phases clear the 45%% trigger: %s go into the 2027 program"
                 % (", ".join(fig["call"]["set"][:-1]) + " and " + fig["call"]["set"][-1]), fontsize=10.5, loc="left")
    f_.text(0.01, 0.01, "Share of each phase's luminaires on their original driver on 1 January 2027 expected to "
            "fail before year 12, under each driver maker's Weibull; labels to the nearest 5%.", fontsize=7.5,
            color="#56616c")
    ax.grid(axis="y", color="#e3e7ea", zorder=0)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    from matplotlib.patches import Patch
    ax.legend(handles=[Patch(color="#1d4e5a", label="In the 2027 program"),
                       Patch(color="#a9b6bf", label="Stays on spot replacement")],
              loc="upper left", frameon=False, fontsize=8.5)
    f_.tight_layout(rect=(0, 0.04, 1, 1))
    f_.savefig(path, metadata={"Software": None})
    plt.close(f_)


def write(target, golden):
    D, G, fig = compute(target)
    os.makedirs(golden, exist_ok=True)
    write_xlsx(os.path.join(golden, P.GOLDEN_XLSX), fig)
    write_png(os.path.join(golden, P.GOLDEN_PNG), fig)
    return fig, dict(D=D, G=G)


def write_figures(path, fig):
    with open(path, "w") as fh:
        json.dump(fig, fh, indent=1, sort_keys=True, default=float)
        fh.write("\n")
