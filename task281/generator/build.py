"""Build task281: the evidence pack into <task>/target, the goldens into <task>/golden, metadata.json and
generator/figures.json, then the checks and the independent verifier, then a second build into a scratch folder
inside the task (removed after) compared byte for byte, then target.zip cut from the final build.

    python3 task281/generator/build.py              everything
    python3 task281/generator/build.py --out DIR    the pack only, into DIR
    python3 task281/generator/build.py --golden     the pack and the goldens, no checks"""
import os
for _v in ("OPENBLAS_NUM_THREADS", "OMP_NUM_THREADS", "MKL_NUM_THREADS"):
    os.environ.setdefault(_v, "1")
import argparse      # noqa: E402
import calendar      # noqa: E402
import hashlib       # noqa: E402
import json          # noqa: E402
import shutil        # noqa: E402
import sys           # noqa: E402
import warnings      # noqa: E402

import numpy as np   # noqa: E402

warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import docs as DC          # noqa: E402
import params as P         # noqa: E402
import records as R        # noqa: E402
import world as Wd         # noqa: E402

TASK = os.path.dirname(HERE)
MTIME = calendar.timegm(P.ASSEMBLED.timetuple())


def write_install_xlsx(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    for y, (head, rows) in sheets.items():
        ws = wb.create_sheet(str(y))
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
        for row in rows:
            ws.append([None if v == "" else v for v in row])
        if y <= 2022:
            for (cell,) in ws.iter_rows(min_row=2, min_col=1, max_col=1):
                cell.number_format = "m/d/yyyy"
        widths = [11, 7, 7, 14, 12, 18, 18, 16, 11, 18, 22]
        for i, wdt in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = wdt
        ws.freeze_panes = "A2"
    wb.save(path)


def build(out, seed=P.SEED, keep=False):
    if os.path.exists(out):
        shutil.rmtree(out)
    os.makedirs(out)
    p = lambda k: os.path.join(out, P.FILES[k])
    w = Wd.World(seed)
    rc = R.Records(w)
    n = {}
    for key, fn in (("register", rc.register_rows), ("wo_legacy", rc.legacy_rows), ("wo_new", rc.new_rows),
                    ("reports", rc.report_rows), ("nlc", rc.nlc_rows), ("billing", rc.billing_rows),
                    ("weather", rc.weather_rows), ("vegetation", rc.vegetation_rows)):
        head, rows = fn()
        R.write_csv(p(key), head, rows)
        n[key] = len(rows)
    sheets = rc.install_sheets()
    write_install_xlsx(p("install"), sheets)
    n["install"] = sum(len(v[1]) for v in sheets.values())
    DC.write_note(p("note"))
    DC.write_coc(p("coc"), w.lots, np.random.default_rng(np.random.SeedSequence(seed + 11)))
    DC.write_tariff(p("tariff"))
    n["coc"] = len(w.lots)
    return dict(w=w if keep else None, rc=rc if keep else None, counts=n)


PDF_DATES = {"note": ("D:20270106143000-05'00'", "D:20270106151200-05'00'"),
             "coc": ("D:20251002091500-04'00'", "D:20251002091500-04'00'"),
             "tariff": ("D:20251118100000-05'00'", "D:20251118100000-05'00'")}
OOXML = {"install": ("Brayle Electric", "Ebony Rivera", "2021-01-12T22:05:00Z", "2025-10-03T19:44:00Z",
                     "TP&L LED conversion install logs", True)}


def finish(out):
    """The containers as their writers save them, then every file's modified time set to the pack's assembly."""
    import postprocess as PP
    for k, (c, m) in PDF_DATES.items():
        PP.normalize_pdf(os.path.join(out, P.FILES[k]), c.encode(), m.encode())
    PP.COMPANY = "Brayle Electric, Inc."
    for k, (cr, lb, c, m, title, excel) in OOXML.items():
        PP.normalize_ooxml(os.path.join(out, P.FILES[k]), cr, lb, c, m, title=title, excel=excel)
    for f in os.listdir(out):
        os.utime(os.path.join(out, f), (MTIME, MTIME))


def finish_golden(golden):
    import postprocess as PP
    PP.COMPANY = P.UTILITY
    PP.normalize_ooxml(os.path.join(golden, P.GOLDEN_XLSX), P.PEOPLE["asset_engineer"], P.PEOPLE["asset_engineer"],
                       "2027-01-12T14:10:00Z", "2027-01-15T21:32:00Z", title="2027 group driver replacement",
                       excel=True)
    for f in os.listdir(golden):
        os.utime(os.path.join(golden, f), (MTIME, MTIME))


SOURCES = {
    "note": "the utility's street lighting lifecycle model note, revision 3",
    "register": "the utility's GIS streetlight inventory, exported 31 December 2026",
    "install": "the installer's daily installation logs, 2021 to 2025",
    "coc": "the luminaire maker's certificates of conformance, one page a delivery lot",
    "wo_legacy": "the legacy work-management system's export at its June 2023 cutover",
    "wo_new": "the enterprise asset management system's work orders, July 2023 to December 2026",
    "reports": "the customer portal's and call centre's outage reports, December 2025 to December 2026",
    "tariff": "the utility's street lighting tariff, Schedule SL",
    "nlc": "Corrinvale's networked lighting control system, monthly node records for 2026",
    "billing": "December 2026 street lighting billing lines by municipal account",
    "weather": "the utility's operations centre weather station, daily",
    "vegetation": "vegetation clearance patrols around streetlights, 2025 and 2026",
}


def metadata(fig, counts):
    src = ("constructed by the task generator (seeded); %s, %s, %s, the towns, people and records are fictional; "
           "the records follow the forms of US utility GIS inventories, installer logs, work-management exports, "
           "customer outage reports and manufacturers' certificates of conformance" % (P.UTILITY, P.MAKER, P.INSTALLER))
    return {
        "task": "task281",
        "title": "%s: the 2027 group driver replacement program across ten LED conversion phases" % P.UTILITY,
        "domain": "General Statistics & Data Science (missing data, censoring and measurement error)",
        "objective": "Descriptive & Distribution Analysis",
        "distractor_files": ["target/" + P.FILES[k] for k in P.DISTRACTORS],
        "deliverables": list(P.DELIVERABLES),
        "provenance": {P.FILES[k]: {"source": "%s; %s" % (SOURCES[k], src), "date": "2027-01-14",
                                    "licence": "CC0 1.0 (constructed for this task)"} for k in sorted(P.FILES)},
        "rows": {P.FILES[k]: v for k, v in sorted(counts.items())},
        "answer": fig["call"]["sentence"],
    }


def digest(folder):
    h = hashlib.sha256()
    for root, _, files in sorted(os.walk(folder)):
        for f in sorted(files):
            fp = os.path.join(root, f)
            h.update(os.path.relpath(fp, folder).encode())
            h.update(open(fp, "rb").read())
    return h.hexdigest()


def main_build(target, golden, seed=P.SEED):
    import deliver as DV
    b = build(target, seed, keep=True)
    finish(target)
    if os.path.exists(golden):
        shutil.rmtree(golden)
    fig, res = DV.write(target, golden)
    finish_golden(golden)
    DV.write_figures(os.path.join(HERE, "figures.json"), fig)
    import writeup as WU
    WU.write(json.load(open(os.path.join(HERE, "figures.json"))))
    with open(os.path.join(TASK, "metadata.json"), "w") as f:
        json.dump(metadata(fig, b["counts"]), f, indent=1, sort_keys=True)
        f.write("\n")
    return b, fig, res


def make_zip(target, zpath):
    import zipfile
    files = sorted(os.listdir(target))
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for f in files:
            zi = zipfile.ZipInfo("target/" + f, date_time=P.ASSEMBLED.timetuple()[:6])
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            z.writestr(zi, open(os.path.join(target, f), "rb").read())


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out")
    ap.add_argument("--golden", action="store_true")
    a = ap.parse_args()
    if a.out:
        build(a.out)
        finish(a.out)
        print("pack written to", a.out)
        return
    target, golden = os.path.join(TASK, "target"), os.path.join(TASK, "golden")
    b, fig, res = main_build(target, golden)
    print("pack:", {k: v for k, v in b["counts"].items()})
    if a.golden:
        return
    import checks as CK
    import verify as VF
    ok = CK.run(target, golden, b, fig, res)
    vok = VF.run(target, golden)
    scratch = os.path.join(TASK, ".rebuild_check")
    b2 = build(os.path.join(scratch, "target"))
    finish(os.path.join(scratch, "target"))
    same = digest(target) == digest(os.path.join(scratch, "target"))
    import deliver as DV
    fig2, _ = DV.write(os.path.join(scratch, "target"), os.path.join(scratch, "golden"))
    finish_golden(os.path.join(scratch, "golden"))
    same_g = digest(golden) == digest(os.path.join(scratch, "golden"))
    shutil.rmtree(scratch)
    print("rebuild byte-identical: target %s, golden %s" % (same, same_g))
    if ok and vok and same and same_g:
        make_zip(target, os.path.join(TASK, "target.zip"))
        print("target.zip", hashlib.sha256(open(os.path.join(TASK, "target.zip"), "rb").read()).hexdigest()[:8])
    else:
        print("NOT zipped: checks %s, verify %s, rebuild %s/%s" % (ok, vok, same, same_g))
        sys.exit(1)


if __name__ == "__main__":
    main()
