"""The golden path for task281, worked from the shipped pack alone (nothing here reads the world).

The call: each conversion luminaire's original-driver lifetime from the installer's logs (retrofit date, phase, lot),
the luminaire maker's certificates of conformance (each lot's driver and its maker) and both work-management systems
(driver replacements, knockdowns, warranty swaps and pole removals, dated at completion); a two-parameter Weibull fitted
by maximum likelihood with suspensions to each driver maker's drivers; each phase's expected share, among its luminaires
still on their original driver on 1 January 2027, of drivers failing before the twelfth anniversary of installation,
given survival to their age on that date; a phase is designated when its share is above 45 per cent.

The asks: each town's 2026 energy saving on the conversion's poles against the sodium fixtures they replaced, on the
tariff's energy accounting; and each town's outages repaired in 2026 more than 72 hours after the first report of them
(or after a patrol found them)."""
import calendar
import csv
import datetime as dt
import os
import re
import subprocess
import zoneinfo

import numpy as np
from scipy.optimize import brentq, minimize

import params as P

EASTERN = zoneinfo.ZoneInfo("America/New_York")
OUTAGE_ACT = {"SL-DRV-RPL": "driver", "SL-PC-RPL": "photocontrol", "SL-WIRING": "wiring", "SL-FUSE-RPL": "fuse",
              "SL-LAMP-RPL": "lamp", "SL-BALLAST-RPL": "ballast", "SL-KNOCKDOWN": "knockdown",
              "SL-LUM-RPL": "luminaire", "SL-UG-CABLE": "ug_cable"}
PLANNED_ACT = {"SL-POLE-REM": "pole_removed", "SL-GLARE": "glare"}
LEG_RES = {"DRVR": "driver", "PEC": "photocontrol", "WIRE": "wiring", "FUSE": "fuse", "LAMP": "lamp",
           "BAL": "ballast", "KD": "knockdown", "LUM": "luminaire", "REM": "pole_removed", "GLR": "glare",
           "UGC": "ug_cable"}
EXIT_KINDS = ("knockdown", "luminaire", "pole_removed")


def f(target, key):
    return os.path.join(target, P.FILES[key])


def pole_of(s):
    """A pole number from any of the forms the systems write: 412345, TPL 412345, SL412345, SL-412345, TPL412345,
    Pole 412345, or an address with the tag after it."""
    m = re.findall(r"(\d{6})", s or "")
    return int(m[-1]) if m else None


# ---------------------------------------------------------------------------------------------- readers
def read_install(target):
    from openpyxl import load_workbook
    wb = load_workbook(f(target, "install"), read_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        head = [str(h) for h in next(rows)]
        ix = {h: i for i, h in enumerate(head)}
        old = "Pole Tag" in ix
        for row in rows:
            if row[0] is None:
                continue
            if old:
                d = row[ix["Date"]].date() if isinstance(row[ix["Date"]], dt.datetime) else row[ix["Date"]]
                rec = dict(pole=pole_of(str(row[ix["Pole Tag"]])), date=d, phase=row[ix["Phase"]],
                           town=row[ix["Municipality"]], serial=row[ix["Luminaire S/N"]],
                           catalog=row[ix["Catalog No."]], lot=row[ix["Lot #"]], removed=row[ix["Removed"]])
            else:
                d = dt.datetime.strptime(row[ix["install_date"]], "%m/%d/%Y").date()
                rec = dict(pole=pole_of(row[ix["pole_id"]]), date=d, phase=row[ix["phase"]], town=row[ix["town"]],
                           serial=row[ix["serial"]], catalog=row[ix["catalog"]], lot=row[ix["lot"]],
                           removed=row[ix["removed_fixture"]])
            assert rec["pole"] not in out, rec
            out[rec["pole"]] = rec
    return out


def read_coc(target):
    """Each lot's LED driver part and its maker, from the certificates of conformance."""
    txt = subprocess.run(["pdftotext", "-layout", f(target, "coc"), "-"], capture_output=True, text=True,
                         check=True).stdout
    lots = {}
    for page in txt.split("\f"):
        m = re.search(r"Lot (HL\d{4}-\d{2})\.", page)
        if not m:
            continue
        d = re.search(r"LED driver, programmable, 0-10 V dimming\s+(\S+)\s+(.+?)\s{2,}(\S+)\s*$", page, re.M)
        assert d, m.group(1)
        assert m.group(1) not in lots
        lots[m.group(1)] = dict(part=d.group(1), maker=d.group(2).strip(), rev=d.group(3))
    return lots


def _ts(s, fmt):
    return dt.datetime.strptime(s, fmt) if s else None


def _wall(s):
    """An offset-stamped local time read as its local wall clock."""
    return dt.datetime.fromisoformat(s).replace(tzinfo=None) if s else None


def read_wos(target):
    """Both systems' work orders as one list: pole, kind, opened, completed (None while open), source, parts."""
    out = []
    with open(f(target, "wo_legacy"), newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            out.append(dict(sys="legacy", wo=r["WO_NO"], pole=pole_of(r["LOCATION"]), kind=LEG_RES.get(r["RES"]),
                            opened=_ts(r["OPENED"], "%m/%d/%Y %H:%M"), reported=None,
                            completed=_ts(r["CLOSED"], "%m/%d/%Y %H:%M"), source="", parts=""))
    with open(f(target, "wo_new"), newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            kind = OUTAGE_ACT.get(r["activity"]) or PLANNED_ACT.get(r["activity"])
            out.append(dict(sys="eam", wo=r["work_order"], pole=pole_of(r["asset_id"]), kind=kind,
                            opened=_wall(r["created_at"]), reported=_wall(r["reported_at"]),
                            completed=_wall(r["completed_at"]) if r["status"] == "CLOSED" else None,
                            source=r["source"], parts=r["parts"]))
    return out


def read_register(target):
    with open(f(target, "register"), newline="", encoding="utf-8") as fh:
        return {pole_of(r["pole_id"]): r for r in csv.DictReader(fh)}


def read_reports(target, utc_as_local=False):
    out = []
    with open(f(target, "reports"), newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            if r["channel"] == "PORTAL":
                u = dt.datetime.strptime(r["received"], "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=dt.timezone.utc)
                t = u.replace(tzinfo=None) if utc_as_local else u.astimezone(EASTERN).replace(tzinfo=None)
            else:
                t = dt.datetime.strptime(r["received"][:-4], "%m/%d/%Y %I:%M %p")
            out.append(dict(id=r["report_id"], ch=r["channel"], t=t, pole=pole_of(r["pole"]),
                            address=r["address"], town=r["town"]))
    return out


def read_nlc(target):
    out = {}
    with open(f(target, "nlc"), newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            out[(pole_of(r["pole_id"]), int(r["month"][5:7]))] = (float(r["dimmed_hours"]), float(r["dim_level_pct"]))
    return out


def read_tariff(target):
    """Input watts by fixture and the monthly burning hours, from Schedule SL."""
    txt = subprocess.run(["pdftotext", "-layout", f(target, "tariff"), "-"], capture_output=True, text=True,
                         check=True).stdout
    watts = {}
    for m in re.finditer(r"^\s*(LED|HPS|MH) (\d+) W[^\n]*?\s(\d+)\s+\$", txt, re.M):
        watts[(m.group(1), int(m.group(2)))] = int(m.group(3))
    hrs = re.search(r"Hours\s+((?:\d+\s+){12})", txt).group(1).split()
    return watts, [int(h) for h in hrs]


def load(target):
    return dict(inst=read_install(target), coc=read_coc(target), wos=read_wos(target), reg=read_register(target),
                reports=read_reports(target), nlc=read_nlc(target), tariff=read_tariff(target))


# ---------------------------------------------------------------------------------------------- lifetimes
def lifetimes(D, failure_kinds=("driver",), date_field="completed", exits=EXIT_KINDS):
    """One row per conversion luminaire: years to its original driver's failure or suspension, the failure flag,
    whether it is still on its original driver when the record closes, its age on 1 January 2027, phase and maker."""
    end = P.RECORD_END
    by_pole = {}
    for w in D["wos"]:
        if w["pole"] in D["inst"]:
            by_pole.setdefault(w["pole"], []).append(w)
    rows = []
    for pole, I in D["inst"].items():
        t0 = I["date"]
        fail = exit_ = None
        for w in by_pole.get(pole, []):
            d = w[date_field]
            if d is None:
                continue
            d = d.date()
            if d <= t0 or d > end:
                continue
            if w["kind"] in failure_kinds and (fail is None or d < fail):
                fail = d
            if w["kind"] in exits and (exit_ is None or d < exit_):
                exit_ = d
        failed = fail is not None and (exit_ is None or fail <= exit_)
        stop = fail if failed else min(exit_ or end, end)
        rows.append(dict(pole=pole, phase=I["phase"], maker=D["coc"][I["lot"]]["maker"], lot=I["lot"],
                         watts=int(re.search(r"-(\d+)W-", I["catalog"]).group(1)),
                         t=(stop - t0).days / P.YEAR_D, d=failed,
                         alive=(not failed) and (exit_ is None),
                         age=(P.AS_OF - t0).days / P.YEAR_D, age_close=(end - t0).days / P.YEAR_D))
    return rows


def weibull_mle(t, d):
    """Two-parameter Weibull by maximum likelihood with suspensions: the shape solves the profile score equation,
    the scale follows in closed form."""
    t = np.maximum(np.asarray(t, float), 1e-9)
    d = np.asarray(d, bool)
    r = d.sum()
    if r == 0:
        return 1.0, np.inf                      # no failure on record: no finite fit
    lt = np.log(t)
    mean_lf = lt[d].mean()

    def score(k):
        tk = t ** k
        return (tk * lt).sum() / tk.sum() - 1.0 / k - mean_lf

    k = brentq(score, 0.05, 30.0, xtol=1e-14)
    lam = ((t ** k).sum() / r) ** (1.0 / k)
    return k, lam


def F(t, k, lam):
    return 1.0 - np.exp(-(np.asarray(t, float) / lam) ** k)


def cond_fail(a, k, lam, h=P.HORIZON_Y):
    """Probability of a failure before age h given survival to age a."""
    a = np.asarray(a, float)
    return (F(h, k, lam) - F(a, k, lam)) / (1.0 - F(a, k, lam))


def phase_shares(rows, prob):
    """Each phase's expected share of its original-driver luminaires failing before the horizon; prob(row) gives a
    luminaire's probability."""
    out = {}
    for pc in P.PHASE_CODES:
        rs = [r for r in rows if r["phase"] == pc and r["alive"]]
        out[pc] = dict(n=len(rs), share=float(np.mean([prob(r) for r in rs])))
    return out


def golden_call(D, rows=None):
    rows = rows if rows is not None else lifetimes(D)
    fits = {}
    for mk in sorted({r["maker"] for r in rows}):
        rr = [r for r in rows if r["maker"] == mk]
        k, lam = weibull_mle([r["t"] for r in rr], [r["d"] for r in rr])
        fits[mk] = dict(k=k, lam=lam, n=len(rr), failures=int(sum(r["d"] for r in rr)),
                        F12=float(F(P.HORIZON_Y, k, lam)))
    sh = phase_shares(rows, lambda r: cond_fail(r["age"], fits[r["maker"]]["k"], fits[r["maker"]]["lam"]))
    for pc, v in sh.items():
        rs = [r for r in rows if r["phase"] == pc and r["alive"]]
        v["by_maker"] = {mk: sum(1 for r in rs if r["maker"] == mk) for mk in fits}
        v["expected"] = v["share"] * v["n"]
        v["designated"] = v["share"] > P.TRIGGER
        v["failed_to_date"] = sum(1 for r in rows if r["phase"] == pc and r["d"])
        v["installed"] = sum(1 for r in rows if r["phase"] == pc)
    return dict(fits=fits, phases=sh, set=[pc for pc in P.PHASE_CODES if sh[pc]["designated"]], rows=rows)


# ---------------------------------------------------------------------------------------------- the rungs
def rung_pooled(rows, cond=True):
    k, lam = weibull_mle([r["t"] for r in rows], [r["d"] for r in rows])
    if cond:
        return phase_shares(rows, lambda r: cond_fail(r["age"], k, lam)), (k, lam)
    return phase_shares(rows, lambda r: float(F(P.HORIZON_Y, k, lam))), (k, lam)


def weibull_groups(rows, key, shared_shape=True):
    """A Weibull with one scale per group of `key` and one shared shape (or a shape per group)."""
    groups = sorted({key(r) for r in rows})
    t = np.maximum(np.array([r["t"] for r in rows]), 1e-9)
    d = np.array([r["d"] for r in rows])
    g = np.array([groups.index(key(r)) for r in rows])
    if not shared_shape:
        return {gg: weibull_mle(t[g == i], d[g == i]) for i, gg in enumerate(groups)}
    lt = np.log(t)

    def prof(k):
        ll = 0.0
        for i in range(len(groups)):
            m = g == i
            r = d[m].sum()
            s = (t[m] ** k).sum()
            if r == 0:
                continue
            lam_k = s / r
            ll += r * np.log(k) - r * np.log(lam_k) + (k - 1) * lt[m][d[m]].sum() - r
        return -ll

    res = minimize(lambda x: prof(np.exp(x[0])), [np.log(3.0)], method="Nelder-Mead",
                   options={"xatol": 1e-12, "fatol": 1e-12, "maxiter": 4000})
    k = float(np.exp(res.x[0]))
    out = {}
    for i, gg in enumerate(groups):
        m = g == i
        r = d[m].sum()
        out[gg] = (k, ((t[m] ** k).sum() / r) ** (1 / k) if r else np.inf)
    return out


def rung_groups(rows, key, shared_shape=True):
    fits = weibull_groups(rows, key, shared_shape)
    return phase_shares(rows, lambda r: cond_fail(r["age"], *fits[key(r)]) if np.isfinite(fits[key(r)][1]) else 0.0)


def lfp_fit(t, d):
    """Limited-failure-population Weibull: a fraction p susceptible, the rest never failing."""
    t = np.maximum(np.asarray(t, float), 1e-9)
    d = np.asarray(d, bool)

    def nll(x):
        k, lam, p = np.exp(x[0]), np.exp(x[1]), 1 / (1 + np.exp(-x[2]))
        z = (t / lam) ** k
        lf = np.log(p) + np.log(k / lam) + (k - 1) * np.log(t / lam) - z
        ls = np.log(1 - p + p * np.exp(-z))
        return -(lf[d].sum() + ls[~d].sum())

    best = None
    for k0, l0, p0 in ((3, 10, 0.5), (2, 20, 0.8), (4, 8, 0.3), (3, 12, 0.95)):
        r = minimize(nll, [np.log(k0), np.log(l0), np.log(p0 / (1 - p0))], method="Nelder-Mead",
                     options={"xatol": 1e-10, "fatol": 1e-10, "maxiter": 20000})
        if best is None or r.fun < best.fun:
            best = r
    k, lam, p = np.exp(best.x[0]), np.exp(best.x[1]), 1 / (1 + np.exp(-best.x[2]))
    return k, lam, p


def rung_lfp(rows):
    k, lam, p = lfp_fit([r["t"] for r in rows], [r["d"] for r in rows])

    def prob(r):
        Sa = 1 - p + p * np.exp(-(r["age"] / lam) ** k)
        S12 = 1 - p + p * np.exp(-(P.HORIZON_Y / lam) ** k)
        return (Sa - S12) / Sa
    return phase_shares(rows, prob), (k, lam, p)


def rung_lfp_phase(rows):
    """A shared Weibull with a susceptible fraction per phase (the cure model with the phase as its covariate)."""
    t = np.maximum(np.array([r["t"] for r in rows]), 1e-9)
    d = np.array([r["d"] for r in rows])
    g = np.array([P.PHASE_CODES.index(r["phase"]) for r in rows])

    def nll(x):
        k, lam = np.exp(x[0]), np.exp(x[1])
        p = 1 / (1 + np.exp(-x[2:]))[g]
        z = (t / lam) ** k
        lf = np.log(p) + np.log(k / lam) + (k - 1) * np.log(t / lam) - z
        ls = np.log(1 - p + p * np.exp(-z))
        return -(lf[d].sum() + ls[~d].sum())

    best = None
    for p0 in (0.3, 0.6, 0.9):
        r = minimize(nll, np.r_[np.log(3.2), np.log(10.0), np.full(len(P.PHASE_CODES), np.log(p0 / (1 - p0)))],
                     method="L-BFGS-B")
        if best is None or r.fun < best.fun:
            best = r
    k, lam = np.exp(best.x[:2])
    pp = 1 / (1 + np.exp(-best.x[2:]))
    out = {}
    for i, pc in enumerate(P.PHASE_CODES):
        a = np.array([r["age"] for r in rows if r["phase"] == pc and r["alive"]])
        Sa = 1 - pp[i] + pp[i] * np.exp(-(a / lam) ** k)
        S12 = 1 - pp[i] + pp[i] * np.exp(-(P.HORIZON_Y / lam) ** k)
        out[pc] = dict(share=float(np.mean((Sa - S12) / Sa)), p=float(pp[i]))
    return out


def rung_maker_uncond(D, G):
    """The at-risk maker's year-twelve failure probability times the phase's share of its drivers, unconditioned."""
    rows = G["rows"]
    out = {}
    for pc in P.PHASE_CODES:
        rs = [r for r in rows if r["phase"] == pc]
        out[pc] = dict(share=float(np.mean([G["fits"][r["maker"]]["F12"] for r in rs])))
    return out


def rung_zero_cure(G, cure_maker):
    rows = G["rows"]
    fits = G["fits"]
    return phase_shares(rows, lambda r: 0.0 if r["maker"] == cure_maker
                        else cond_fail(r["age"], fits[r["maker"]]["k"], fits[r["maker"]]["lam"]))


def rung_naive(rows):
    return {pc: dict(share=float(np.mean([r["d"] for r in rows if r["phase"] == pc]))) for pc in P.PHASE_CODES}


# ---------------------------------------------------------------------------------------------- ask E: energy
def _days(y, m):
    return calendar.monthrange(y, m)[1]


def pole_timeline(D, pole, I):
    """The pole's 2026 service: (from_date, LED nominal watts) steps, and the day it was removed (or None)."""
    evs = sorted([w for w in D["wos"] if w["pole"] == pole and w["completed"] is not None
                  and w["completed"].date() > I["date"]], key=lambda w: w["completed"])
    watts = int(re.search(r"-(\d+)W-", I["catalog"]).group(1))
    steps = [(I["date"], watts)]
    removed = None
    for w in evs:
        d = w["completed"].date()
        if w["kind"] == "glare" and not (w["parts"] or "").startswith("HSS"):
            steps.append((d, int(D["reg"][pole]["lamp_watts"])))
        elif w["kind"] in ("knockdown", "luminaire"):
            m = re.search(r"HRX-\d-(\d+)W-", w["parts"] or "")
            if m:
                steps.append((d, int(m.group(1))))
        elif w["kind"] == "pole_removed":
            removed = d if removed is None else min(removed, d)
    return steps, removed


def energy_saving(D, removal_day_in=False, change_day_new=True, impute_missing=False):
    """Each town's 2026 saving in MWh on the converted poles: the removed sodium fixture's input watts against the LED's,
    for the tariff's burning hours in proportion to the days in service, the LED at its dimmed input for the dimmed hours
    its node recorded."""
    watts_tab, burn = D["tariff"]
    by_pole = {}
    for w in D["wos"]:
        by_pole.setdefault(w["pole"], []).append(w)
    res = {}
    per_pole = {}
    for pole, I in D["inst"].items():
        hm = re.match(r"HPS (\d+)W (\w+)", I["removed"])
        hps_in = watts_tab[("HPS", int(hm.group(1)))]
        sub = dict(D, wos=by_pole.get(pole, []))
        steps, removed = pole_timeline(sub, pole, I)
        if removed is not None and removed < dt.date(2026, 1, 1):
            continue
        total = 0.0
        for m in range(1, 13):
            nd = _days(2026, m)
            days = [dt.date(2026, m, x) for x in range(1, nd + 1)]
            if removed is not None:
                days = [x for x in days if (x <= removed if removal_day_in else x < removed)]
            if not days:
                continue
            per_day_h = burn[m - 1] / nd
            led_w = []
            for x in days:
                cur = steps[0][1]
                for s_d, s_w in steps:
                    if (s_d <= x) if change_day_new else (s_d < x):
                        cur = s_w
                led_w.append(watts_tab[("LED", cur)])
            led_w = np.array(led_w, float)
            h = per_day_h * len(days)
            e_hps = hps_in * h
            e_led = (led_w * per_day_h).sum()
            rec = D["nlc"].get((pole, m))
            if rec is None and impute_missing:
                rec = _impute(D, pole, m)
            if rec is not None:
                dh, lvl = rec
                # the dimmed hours spread over the month's days in service at the input of each day
                e_led -= (led_w.mean()) * (1 - lvl / 100.0) * dh
            total += e_hps - e_led
        per_pole[pole] = total
        res[I["town"]] = res.get(I["town"], 0.0) + total
    return {t: v / 1e6 for t, v in res.items()}, per_pole


def _impute(D, pole, m):
    vals = [D["nlc"][(pole, mm)] for mm in range(1, 13) if (pole, mm) in D["nlc"]]
    if not vals:
        return None
    return (float(np.mean([v[0] for v in vals])) * P.BURN_H[m - 1] / np.mean([P.BURN_H[mm - 1] for mm in range(1, 13)
                                                                              if (pole, mm) in D["nlc"]]),
            vals[0][1])


# ---------------------------------------------------------------------------------------------- ask R: repairs
def pole_town(D):
    t = {p: r["town"] for p, r in D["reg"].items()}
    for p, I in D["inst"].items():
        t.setdefault(p, I["town"])
    return t


def report_poles(D):
    """Each report's pole: the number it carries, or its address looked up in the register within its town."""
    addr = {}
    for p, r in D["reg"].items():
        addr.setdefault((r["town"], r["nearest_address"]), []).append(p)
    out = []
    for rp in D["reports"]:
        p = rp["pole"]
        if p is None:
            c = addr.get((rp["town"], rp["address"]), [])
            p = c[0] if len(c) == 1 else None
        out.append(dict(rp, pole_res=p))
    return out


def repairs(D, first="all_reports", threshold_h=72.0, include_planned=False, include_patrol=True, strict=True,
            circuits=True, lookback_days=None):
    """Each town's outages repaired in 2026, and how many were repaired more than 72 hours after the first report of
    them, or, for an outage the Company found before any report, after it was found. An underground cable fault
    darkens every light on its circuit, so a report about any of them is a report of the fault."""
    town = pole_town(D)
    reps = report_poles(D)
    feed = {p: r["feed"] for p, r in D["reg"].items()}
    circ = {}
    for p, fd in feed.items():
        if fd.startswith("UG"):
            circ.setdefault(fd, []).append(p)
    by_pole_rep = {}
    for rp in reps:
        if rp["pole_res"] is not None:
            by_pole_rep.setdefault(rp["pole_res"], []).append(rp)
    wos = [w for w in D["wos"] if w["sys"] == "eam" and w["completed"] is not None]

    def lights(w):
        if circuits and w["kind"] == "ug_cable":
            return circ.get(feed.get(w["pole"]), [w["pole"]])
        return [w["pole"]]

    cover = {}
    for w in wos:
        for p in lights(w):
            cover.setdefault(p, []).append(w["completed"])
    kinds = set(OUTAGE_ACT.values()) | (set(PLANNED_ACT.values()) if include_planned else set())
    total, late = {}, {}
    detail = []
    for w in wos:
        if w["kind"] not in kinds or w["completed"].year != 2026:
            continue
        if not include_patrol and w["source"] in ("PATROL", "TROUBLE") and w["kind"] != "ug_cable":
            continue
        ls = lights(w)
        prev = [c for p in ls for c in cover.get(p, []) if c < w["completed"]]
        prev = max(prev) if prev else None
        if lookback_days is not None:
            prev = w["opened"] - dt.timedelta(days=lookback_days)
        cand = [rp["t"] for p in ls for rp in by_pole_rep.get(p, [])
                if (prev is None or rp["t"] > prev) and rp["t"] <= w["completed"]]
        fallback = w["reported"] if w["source"].startswith("CR-") else w["opened"]
        if first == "reported_at":
            t_first = fallback
        else:
            t_first = min(cand + ([w["reported"]] if w["source"].startswith("CR-") else [])) if cand else fallback
        hrs = (w["completed"] - t_first).total_seconds() / 3600.0
        miss = hrs > threshold_h if strict else hrs >= threshold_h
        tw = town[w["pole"]]
        total[tw] = total.get(tw, 0) + 1
        late[tw] = late.get(tw, 0) + int(miss)
        detail.append(dict(wo=w["wo"], pole=w["pole"], town=tw, hrs=hrs, first=t_first, src=w["source"],
                           kind=w["kind"], n_reports=len(cand)))
    return late, total, detail
