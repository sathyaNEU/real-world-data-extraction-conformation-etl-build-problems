"""An independent check of task458's goldens from the bundle alone, on its own code path: NPPES as the weekly files
update it through 1 November 2026, the practising prescribers, the adult members by client identifier, the final
claims, each county's members with each condition (each identified member weighted by 1 over the share of the members
enrolled all year with the condition whom a claim identifies within the same number of months), the three standards,
the designation, and the fee-for-service payments by county and quarter and by county and group; each compared with
the deliverables."""
import csv
import datetime as dt
import os
import sys

import pandas as pd

BH = {"2084P0800X", "2084P0804X", "2084P0802X", "2084P0805X", "2084A0401X", "207QA0401X", "207RA0401X", "363LP0808X",
      "364SP0808X"}
VISIT = {"99202", "99203", "99204", "99205", "99211", "99212", "99213", "99214", "99215", "90791", "90792"}
NAMES = {"53001": "Adams", "53005": "Benton", "53007": "Chelan", "53017": "Douglas", "53021": "Franklin", "53025": "Grant",
         "53037": "Kittitas", "53047": "Okanogan", "53071": "Walla Walla", "53077": "Yakima"}


def months_between(a, b):
    out = set()
    y, m = a.year, a.month
    while (y, m) <= (b.year, b.month):
        if (2025, 7) <= (y, m) <= (2026, 6):
            out.add((y, m))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


def main(target, golden):
    t = lambda f: os.path.join(target, f)
    # NPPES on 1 November: the monthly file, then each weekly file ending by 1 November, record by record
    files = ["npidata_pfile_wa_20050523-20261011.csv", "npidata_pfile_wa_20261012-20261018.csv",
             "npidata_pfile_wa_20261019-20261025.csv", "npidata_pfile_wa_20261026-20261101.csv"]
    rec, live = {}, {}
    for f in files:
        for r in pd.read_csv(t(f), dtype=str, keep_default_na=False).to_dict("records"):
            rec[r["NPI"]] = r
            if r["Entity Type Code"] == "1" or r["Entity Type Code"] == "2":
                live[r["NPI"]] = r
    z = pd.read_excel(t("ZIP_COUNTY_062026.xlsx"), dtype={"ZIP": str, "COUNTY": str})
    zc = z.sort_values(["ZIP", "BUS_RATIO"], ascending=[True, False]).drop_duplicates("ZIP").set_index("ZIP")["COUNTY"]

    def prim(r):
        for k in range(1, 5):
            if r["Healthcare Provider Primary Taxonomy Switch_%d" % k] == "Y":
                return r["Healthcare Provider Taxonomy Code_%d" % k]
        return ""
    bh, place = set(), {}
    for npi, r in rec.items():
        src = r if r["Entity Type Code"] else live.get(npi)
        if src is None or src["Entity Type Code"] != "1" or prim(src) not in BH:
            continue
        bh.add(npi)
        if r["Entity Type Code"] and r["Provider Business Practice Location Address State Name"] == "WA":
            place[npi] = zc.get(r["Provider Business Practice Location Address Postal Code"][:5])
    # members
    s = pd.read_csv(t("enrollment_spans_2025-07-01_2026-06-30.csv"), dtype=str, keep_default_na=False)
    s["a"] = pd.to_datetime(s["span_start_dt"]).dt.date.map(lambda d: max(d, dt.date(2025, 7, 1)))
    s["b"] = s["span_end_dt"].map(lambda x: dt.date.fromisoformat(x) if x else dt.date(2026, 6, 30)).map(
        lambda d: min(d, dt.date(2026, 6, 30)))
    s = s[s["a"] <= s["b"]]
    per = {}
    for r in s.to_dict("records"):
        p = per.setdefault(r["client_id"], {"months": set(), "end": None, "fips": None, "grp": None, "start": None})
        p["months"] |= months_between(r["a"], r["b"])
        if p["end"] is None or (r["b"], r["a"]) > (p["end"], p["start"]):
            p["end"], p["start"], p["fips"], p["grp"] = r["b"], r["a"], r["res_county_fips"], r["elig_group_cd"]
    mem = pd.DataFrame([{"cid": k, "county": NAMES[v["fips"]], "grp": v["grp"], "full": len(v["months"]) == 12,
                         "m": len(v["months"])} for k, v in per.items()]).set_index("cid")
    to_cid = dict(zip(s["mbr_id"], s["client_id"]))
    # final claims
    c = pd.read_csv(t("bh_claims_encounters_dos_2025-07-01_2026-06-30.csv"), dtype=str, keep_default_na=False)
    replaced = set(c.loc[c["clm_freq_cd"] != "1", "orig_icn"])
    c = c[(c["clm_freq_cd"] != "8") & ~c["clm_icn"].isin(replaced)].copy()
    c["cid"] = c["mbr_id"].map(to_cid)
    smi = {x.replace(".", "") for x in ("F20.0 F20.1 F20.2 F20.3 F20.5 F20.81 F20.89 F20.9 F21 F22 F25.0 F25.1 F25.8 F25.9 "
                                        "F28 F29 F31.0 F31.10 F31.11 F31.12 F31.13 F31.2 F31.30 F31.31 F31.32 F31.4 F31.5 "
                                        "F31.60 F31.61 F31.62 F31.63 F31.64 F31.70 F31.71 F31.72 F31.73 F31.74 F31.75 "
                                        "F31.76 F31.77 F31.78 F31.81 F31.89 F31.9 F32.3 F33.3").split()}
    dxcols = ["dx_cd_1", "dx_cd_2", "dx_cd_3", "dx_cd_4"]
    has_smi = c[dxcols].isin(smi).any(axis=1)
    has_oud = c[dxcols].apply(lambda col: col.str.match(r"^F11[12]")).any(axis=1)
    mem["smi"] = mem.index.isin(set(c.loc[has_smi, "cid"]))
    mem["oud"] = mem.index.isin(set(c.loc[has_oud, "cid"]))
    month = (c["svc_from_dt"].str[:4].astype(int) - 2025) * 12 + c["svc_from_dt"].str[5:7].astype(int) - 6
    for cond, has in (("smi", has_smi), ("oud", has_oud)):
        first = month[has].groupby(c.loc[has, "cid"]).min()
        f = first.reindex(mem.index[mem["full"] & mem[cond]])
        share = [float((f <= k).mean()) for k in range(1, 13)]
        mem["w_" + cond] = [1.0 / share[m - 1] if hit else 0.0 for m, hit in zip(mem["m"], mem[cond])]
    mem["seen"] = mem.index.isin(set(c.loc[c["proc_cd"].isin(VISIT) & c["rendering_npi"].isin(bh), "cid"]))
    rendering = set(c["rendering_npi"])
    pract = {}
    for npi in bh:
        if npi in place and place[npi] in NAMES and npi in rendering:
            pract[NAMES[place[npi]]] = pract.get(NAMES[place[npi]], 0) + 1
    out = {}
    for county, g in mem.groupby("county"):
        e_smi, e_oud = g["w_smi"].sum(), g["w_oud"].sum()
        seen = int((g["smi"] & g["seen"]).sum())
        n = pract.get(county, 0)
        v = {"n": n, "smi": e_smi, "oud": e_oud, "s1": n / e_smi * 1000, "s2": n / e_oud * 1000, "s3": seen / e_smi}
        v["fails"] = int(v["s1"] < 16.5) + int(v["s2"] < 23.0) + int(v["s3"] < 0.68)
        out[county] = v
    designated = sorted(k for k, v in out.items() if v["fails"] >= 2)
    # payments: final fee-for-service claims rendered by a behavioral-health prescriber, each claim's paid amount once
    ffs = c[c["fin_arrangement_cd"] == "FFS"]
    claims = ffs[ffs["clm_icn"].isin(set(ffs.loc[ffs["rendering_npi"].isin(bh), "clm_icn"]))]
    first = claims.groupby("clm_icn").first()
    first = first.assign(county=first["cid"].map(mem["county"]), grp=first["cid"].map(mem["grp"]),
                         q=first["svc_from_dt"].str[:7].map(lambda ym: "%s Q%d" % (ym[:4], (int(ym[5:7]) - 1) // 3 + 1)),
                         paid=first["paid_amt"].astype(float))
    pay_q = first.groupby(["county", "q"])["paid"].sum()
    pay_g = first.groupby(["county", "grp"])["paid"].sum()
    # the deliverables
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(golden, "prescriber_pilot_county_scorecard_2027.xlsx"), data_only=True)
    ws = wb["Designation 2027"]
    rows = [r for r in ws.iter_rows(values_only=True) if r and r[0] in out]
    problems = []
    for r in rows:
        v = out[r[0]]
        if r[2] != v["n"]:
            problems.append(("prescribers", r[0], r[2], v["n"]))
        if r[3] != int(round(v["smi"] / 10.0)) * 10 or r[4] != int(round(v["oud"] / 10.0)) * 10:
            problems.append(("need", r[0], r[3], r[4], v["smi"], v["oud"]))
        if abs(r[5] - v["s1"]) > 1e-6 or abs(r[7] - v["s2"]) > 1e-6 or abs(r[9] - v["s3"]) > 1e-9:
            problems.append(("standards", r[0]))
        if (r[12] == "Yes") != (v["fails"] >= 2):
            problems.append(("designated", r[0]))
    ws2 = wb["FFS payments"]
    qs = ["2025 Q3", "2025 Q4", "2026 Q1", "2026 Q2"]
    for r in ws2.iter_rows(values_only=True):
        if r and r[0] in out:
            for k, q in enumerate(qs, start=1):
                if r[k] != int(round(pay_q.get((r[0], q), 0.0))):
                    problems.append(("payment", r[0], q, r[k], pay_q.get((r[0], q), 0.0)))
    with open(os.path.join(golden, "county_adult_members_2026_conformed.csv")) as fh:
        for r in csv.DictReader(fh):
            n = int(((mem["county"] == r["county"]) & (mem["grp"] == r["elig_group_cd"])).sum())
            if int(r["adult_members"]) != n or int(r["ffs_paid_bh_prescriber_usd"]) != int(round(pay_g.get((r["county"], r["elig_group_cd"]), 0.0))):
                problems.append(("conformed", r["county"], r["elig_group_cd"]))
    print("verify: designated %s; %d scorecard rows, %d problems" % (", ".join(designated), len(rows), len(problems)))
    for p in problems[:20]:
        print("  ", p)
    if designated != ["Adams", "Douglas", "Franklin", "Grant", "Okanogan"] or problems:
        raise SystemExit("verify failed")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
