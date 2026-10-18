"""Build task173: the evidence pack into <task>/target, the goldens into <task>/golden, metadata.json and
generator/figures.json, then the checks and the independent verifier, then a second full build into a scratch folder
compared byte for byte (and deleted), then target.zip cut from the pack.

    python3 task173/generator/build.py              everything
    python3 task173/generator/build.py --out DIR    the pack only, into DIR"""
from __future__ import annotations

import argparse
import calendar
import datetime as dt
import hashlib
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import acd as A             # noqa: E402
import common as K          # noqa: E402
import des as DS            # noqa: E402
import docs as DOC          # noqa: E402
import postprocess as PP    # noqa: E402
import records as RC        # noqa: E402
import world as W           # noqa: E402

TASK = os.path.dirname(HERE)
SCRATCH = os.environ.get("TASK173_SCRATCH", "/tmp")
PULL_DONE = (2026, 10, 30, 9, 15, 0)                       # the pull's end, Eastern
MTIME = calendar.timegm((2026, 10, 30, 13, 15, 0, 0, 0, 0))  # the same moment, UTC

F_ROSTER = RC.F_ROSTER
F_CSAT = "csat_survey_responses_2026.csv"
F_RETURNS = "finance_live_returns_weekly.csv"


class Built:
    pass


def world():
    B = Built()
    B.people = W.build_people()
    RC.assign_teams(B.people)
    B.recall = RC.recall_status(B.people)
    B.old = {}
    for y in K.SEASONS_OLD:
        v = W.BENCH[y]["vendor"]
        vend = W.vendor_agents(y, v[0]) if v else []
        arr, staff, _meta = W.simulate(y, B.people, vend)
        B.old[y] = (arr, staff)
    v = W.BENCH[K.SEASON_CRM]["vendor"]
    B.vend26 = W.vendor_agents(K.SEASON_CRM, v[0])
    (B.recs, B.atts, B.shifts, B.brks, B.staff, B.meta, B.vev) = DS.simulate_2026(B.people, B.vend26)
    B.delay = RC.connect_seconds(B.recs)
    B.crm, B.ids = RC.crm_rows(B.recs, B.staff, B.delay)
    B.D = A.Directory(B.staff, B.vev)
    B.sessions, B.clocks = A.vendor_timelines(B.recs, B.atts, B.shifts, B.brks, B.staff, B.vev, B.D)
    B.sessions += A.lead_sessions(B.D)
    B.invoices = A.invoice_lines(B.clocks, B.D)
    # whether each 2026 record's case came back after it (for the survey's answers)
    order = np.lexsort((B.recs["start"], B.recs["day"], B.recs["case"]))
    cs = B.recs["case"][order]
    nxt = np.zeros(len(cs), dtype=bool)
    nxt[:-1] = cs[1:] == cs[:-1]
    B.followed = np.zeros(len(cs), dtype=bool)
    B.followed[order] = nxt
    # the 2026 fortnight's new cases as the CRM counts them, and the demand plan's 2027 figure
    f0, f1 = K.fortnight(K.SEASON_CRM)
    n26 = 0
    for r in B.crm:
        if r[2] == "Live Experts EN" and r[4] == "New case":
            d = K.eastern(dt.datetime.strptime(r[6], "%Y-%m-%dT%H:%M:%SZ")).date()
            n26 += f0 <= d <= f1
    B.F26 = n26
    B.F27 = int(round(n26 * (1 + K.GROWTH_2027)))
    return B


def _xlsx(path, sheets, widths=None):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    first = True
    for name, head, rows in sheets:
        ws = wb.active if first else wb.create_sheet(name)
        ws.title = name
        first = False
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append(r)
        for k, col in enumerate(head):
            w = max(len(str(col)), *(len(str(r[k])) if r[k] is not None else 0 for r in rows[:400])) if rows else len(col)
            ws.column_dimensions[openpyxl.utils.get_column_letter(k + 1)].width = min(46, max(9, w + 2))
        ws.freeze_panes = "A2"
    wb.save(path)


def old_history(B):
    """Each old season's deadline-fortnight new tickets on the English queue (old system), and 2026's from the CRM."""
    out = []
    for y in K.SEASONS_OLD:
        arr, staff = B.old[y]
        f0, f1 = K.fortnight(y)
        m = (arr["queue"] == 0) & (arr["pos"] == 1) & (arr["day"] >= f0.toordinal()) & (arr["day"] <= f1.toordinal())
        out.append((y, int(m.sum())))
    out.append((K.SEASON_CRM, B.F26))
    return out


def build_data(out, B):
    p = lambda f: os.path.join(out, f)
    n = {}
    n["crm"] = K.write_csv(p(RC.F_CRM), RC.CRM_COLS, B.crm)
    # roster with the recall, and the outsourcer's 2026 agents
    rows = RC.roster_rows(B.people, B.recall)
    cols = list(rows[0].keys())
    f0, f1 = K.fortnight(K.SEASON_CRM)
    v0 = f0 - dt.timedelta(7)
    vrows = [[e.emp, "Moncton", v0, f1] for e in B.vend26]
    _xlsx(p(F_ROSTER), [("Seasonal experts", cols, [[r[c] for c in cols] for r in rows]),
                        ("HCS agents 2026", ["Agent", "Site", "First Day On Queue", "Last Day On Queue"], vrows)])
    n["roster"] = len(rows)
    hist = RC.history_rows(B.people)
    n["history"] = K.write_csv(p(RC.F_HISTORY), list(hist[0].keys()), hist)
    old_a = [r for y in (2021, 2022) for r in RC.old_rows(y, *B.old[y])]
    old_b = [r for y in (2023, 2024, 2025) for r in RC.old_rows(y, *B.old[y])]
    n["old_a"] = K.write_csv(p(RC.F_OLD_A), RC.OLD_A_COLS, old_a)
    n["old_b"] = K.write_csv(p(RC.F_OLD_B), RC.OLD_B_COLS, old_b)
    # the demand plan
    days, ints = RC.demand_plan_rows(B.F27)
    basis = RC.demand_basis(B.F26, K.GROWTH_2027, old_history(B))
    basis[-1][4] = B.F27
    _xlsx(p(RC.F_DEMAND), [("Daily forecast", ["Date", "Day", "New expert cases"],
                            [[d, K.WEEKDAY[d.weekday()], int(x)] for d, x in zip(days, ints)]),
                           ("Basis", ["Season", "Filing deadline", "Fortnight from", "Fortnight to",
                                      "New expert cases in the fortnight", "Source"], basis)])
    n["demand"] = len(days)
    # the ACD and the outsourcer
    teams = {e.emp: e.team.get(K.SEASON_CRM, "") for e in B.people if K.SEASON_CRM in e.seasons}
    n["users"] = K.write_csv(p(A.F_USERS), A.USER_COLS, A.user_rows(B.D, teams))
    n["states"] = K.write_csv(p(A.F_STATES), A.STATE_COLS, A.state_rows(B.sessions, B.D))
    n["contacts"] = K.write_csv(p(A.F_CONTACTS), A.CONTACT_COLS, A.contact_rows(B.atts, B.recs, B.ids, B.D, B.staff))
    ob, leads = A.onboarding_rows(B.D)
    _xlsx(p(A.F_ONBOARD), [("Agents", A.ONBOARD_COLS, ob), ("Team leads", A.LEAD_COLS, leads)])
    n["onboard"] = len(ob)
    n["csat"] = K.write_csv(p(F_CSAT), RC.CSAT_COLS, RC.csat_rows(B.recs, B.ids, B.followed))
    n["returns"] = K.write_csv(p(F_RETURNS), RC.RETURNS_COLS, RC.returns_rows())
    return n


def normalize(out):
    p = lambda f: os.path.join(out, f)
    P = K.PEOPLE
    PP.normalize_ooxml(p(F_ROSTER), P["wfm"], P["wfm"], "2025-12-01T15:10:00Z", "2026-10-29T20:31:00Z", excel=True,
                       company=K.COMPANY)
    PP.normalize_ooxml(p(RC.F_DEMAND), P["demand"], P["demand"], "2026-10-05T13:44:00Z", "2026-10-16T18:02:00Z",
                       excel=True, company=K.COMPANY)
    PP.normalize_ooxml(p(A.F_ONBOARD), P["vendor"], P["vendor"], "2026-03-30T12:20:00Z", "2026-04-08T17:55:00Z",
                       excel=True, company="Hillsborough Contact Services Ltd.")


def build(out, quiet=False, B=None):
    os.makedirs(out, exist_ok=True)
    B = B or world()
    n = build_data(out, B)
    normalize(out)
    DOC.write_all(out, n, B.invoices)
    for f in os.listdir(out):
        os.utime(os.path.join(out, f), (MTIME, MTIME))
    if not quiet:
        for f in sorted(os.listdir(out)):
            print("  %-46s %10d bytes" % (f, os.path.getsize(os.path.join(out, f))))
        print("  rows:", n, "F26", B.F26, "F27", B.F27)
    return B, n


def digest(folder):
    h = {}
    for root, _, files in os.walk(folder):
        for f in sorted(files):
            fp = os.path.join(root, f)
            h[os.path.relpath(fp, folder)] = hashlib.sha256(open(fp, "rb").read()).hexdigest()
    return h


def cut_zip(root):
    target = os.path.join(root, "target")
    zpath = os.path.join(root, "target.zip")
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(os.listdir(target)):
            zi = zipfile.ZipInfo(f, date_time=PULL_DONE)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            z.writestr(zi, open(os.path.join(target, f), "rb").read())
    return zpath


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out")
    ap.add_argument("--no-rebuild", action="store_true")
    ap.add_argument("--pack-only", action="store_true")
    a = ap.parse_args()
    if a.out:
        build(a.out)
        return
    target = os.path.join(TASK, "target")
    if os.path.isdir(target):
        shutil.rmtree(target)
    B, n = build(target)
    if a.pack_only:
        return
    import deliver
    deliver.main(target, os.path.join(TASK, "golden"), os.path.join(HERE, "figures.json"))
    import metadata
    metadata.write(TASK, n)
    r1 = subprocess.run([sys.executable, os.path.join(HERE, "checks.py")])
    r2 = subprocess.run([sys.executable, os.path.join(HERE, "verify.py"), target])
    same = same_g = same_f = True
    if not a.no_rebuild:
        tmp = tempfile.mkdtemp(prefix="b173_", dir=SCRATCH)
        try:
            build(os.path.join(tmp, "target"), quiet=True)
            deliver.main(os.path.join(tmp, "target"), os.path.join(tmp, "golden"), os.path.join(tmp, "figures.json"))
            same = digest(target) == digest(os.path.join(tmp, "target"))
            same_g = digest(os.path.join(TASK, "golden")) == digest(os.path.join(tmp, "golden"))
            same_f = open(os.path.join(HERE, "figures.json"), "rb").read() == open(os.path.join(tmp, "figures.json"),
                                                                                   "rb").read()
            print("rebuild byte-identical: target %s, golden %s, figures %s" % (same, same_g, same_f))
        finally:
            shutil.rmtree(tmp)
    z = cut_zip(TASK)
    print("zip", z, hashlib.sha256(open(z, "rb").read()).hexdigest())
    if r1.returncode or r2.returncode or not (same and same_g and same_f):
        sys.exit(1)


if __name__ == "__main__":
    main()
