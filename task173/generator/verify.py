"""Independent verifier: reads only the bundle (target/), the goldens and figures.json, on its own code path (pandas,
pypdf and openpyxl), and recomputes the call, the components and every ask figure.

    python3 verify.py <target>"""
import datetime as dt
import json
import os
import re
import sys

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
FAILS = []


def ok(label, cond, detail=""):
    print(("  ok   " if cond else "  FAIL ") + label + " " + str(detail))
    if not cond:
        FAILS.append(label)


def norm(a):
    a = re.sub(r"\s+", " ", str(a).strip().upper())
    long = {"STREET": "ST", "AVENUE": "AVE", "DRIVE": "DR", "LANE": "LN", "COURT": "CT", "TRAIL": "TRL", "ROAD": "RD",
            "BOULEVARD": "BLVD", "CIRCLE": "CIR", "PLACE": "PL", "PARKWAY": "PKWY"}
    p = a.split(" ")
    if p[-1] in long:
        p[-1] = long[p[-1]]
    return " ".join(p)


def main(target):
    figs = json.load(open(os.path.join(HERE, "figures.json")))
    svc = pd.read_csv(os.path.join(target, "gis_service_lines_20260930.csv"), dtype=str)
    svc["A"] = svc["SERVICE_ADDRESS"].map(norm)
    by_addr = dict(zip(svc["A"], svc["SERVICE_ID"]))
    nd = svc["A"].str.replace(r"^(\d+) [NSEW] ", r"\1 ", regex=True)
    cnt = nd.value_counts()
    by_nodir = {a: s for a, s in zip(nd, svc["SERVICE_ID"]) if cnt[a] == 1}

    def match(a):
        a = norm(a)
        if a in by_addr:
            return by_addr[a]
        return by_nodir.get(re.sub(r"^(\d+) [NSEW] ", r"\1 ", a))

    eam = pd.read_excel(os.path.join(target, "eam_workorders_2021_2026.xlsx"), dtype=str)
    leg = pd.read_csv(os.path.join(target, "lwms_workorders_2012_2021.csv"), dtype=str)
    e = eam[(eam.WORK_TYPE == "SERVICE LEAK") & (eam.WORK_CLASS == "REP") & (eam.STATUS == "CLOSED")]
    r1 = pd.DataFrame({"svc": e.ASSET_ID, "d": pd.to_datetime(e.COMPLETED).dt.normalize()})
    l = leg[(leg.WO_TYPE == "SVC LEAK") & (leg.STATUS == "CLOSED")]
    r2 = pd.DataFrame({"svc": l.ADDRESS.map(match), "d": pd.to_datetime(l.COMP_DATE, format="%m/%d/%Y")})
    reps = pd.concat([r1, r2], ignore_index=True)
    w0, w1 = pd.Timestamp("2024-10-01"), pd.Timestamp("2026-09-30")
    rate = ((reps.d >= w0) & (reps.d <= w1)).sum() / 2.0
    ok("record rate", abs(rate - figs["step.record_repairs_per_year"]) < 0.06, rate)

    prem = pd.read_csv(os.path.join(target, "cis_premises_20260930.csv"), dtype=str)
    prem["S"] = [s if isinstance(s, str) and s else match(a) for s, a in zip(prem.SERVICE_ID, prem.SERVICE_ADDRESS)]
    p2s = dict(zip(prem.PREMISE_ID, prem.S))
    me = pd.read_csv(os.path.join(target, "cis_meter_events_2012_2026.csv"), dtype=str)
    sets = me[(me.EVENT_TYPE == "SET") & (me.REASON != "NEW")].copy()
    sets["svc"] = sets.PREMISE_ID.map(p2s)
    sets["d"] = pd.to_datetime(sets.EVENT_DATE)
    W = 120
    xs = sets[sets.d + pd.Timedelta(days=W) <= w1]
    rr = reps.dropna().sort_values(["svc", "d"])
    grp = {k: v.values for k, v in rr.groupby("svc")["d"]}
    hits = 0
    for s, d in zip(xs.svc, xs.d):
        v = grp.get(s)
        if v is not None:
            hits += int(((v > d) & (v <= d + pd.Timedelta(days=W))).sum())
    n = len(xs)
    # base rate outside the windows
    inside = 0
    win = {}
    for s, d in zip(xs.svc, xs.d):
        win.setdefault(s, []).append(d)
    for s, ds in win.items():
        v = grp.get(s)
        if v is None:
            continue
        for x in v:
            if any(d < x <= d + pd.Timedelta(days=W) for d in ds):
                inside += 1
    days = (w1 - pd.Timestamp("2012-01-01")).days + 1
    base_rate = (len(rr) - inside) / (len(svc) * days - n * W)
    exc = (hits - base_rate * W * n) / n
    ok("added per 100 exchanges", round(100 * exc, 1) == figs["component.added_per_100_exchanges"], "%.4f" % (100 * exc))
    exw = ((sets.d >= w0) & (sets.d <= w1)).sum() / 2.0
    base27 = rate - exc * exw
    ok("2027 base", int(round(base27, -2)) == figs["component.base_2027"], "%.1f" % base27)
    ami = pd.read_csv(os.path.join(target, "ami_install_schedule_2027.csv"), dtype=str)
    nx = int((ami.WORK_TYPE == "METER EXCHANGE").sum())
    ok("AMI exchanges", nx == figs["component.ami_exchanges_2027"], nx)
    total = base27 + exc * nx      # every scheduled exchange's added repairs fall inside 2027 (a July finish)
    ok("call", int(250 * round(total / 250.0)) == figs["call.service_leak_repairs_2027"], "%.1f" % total)

    mains = pd.read_csv(os.path.join(target, "gis_water_mains_20260930.csv"), dtype=str)
    mc = dict(zip(mains.MAIN_ID, mains.SVC_CENTER))
    mb = eam[(eam.WORK_TYPE == "MAIN BREAK") & eam.WORK_CLASS.isin(["REP", "REPC"]) & (eam.STATUS == "CLOSED")].copy()
    mb["c"] = pd.to_datetime(mb.COMPLETED)
    br = mb.groupby(["ASSET_ID", "REPORTED", "ADDRESS"])["c"].max().reset_index()
    br = br[(br.c >= "2025-10-01") & (br.c <= "2026-09-30")]
    br["cen"] = br.ASSET_ID.map(mc)
    mon = {1: "jan", 2: "feb", 3: "mar", 4: "apr", 5: "may", 6: "jun", 7: "jul", 8: "aug", 9: "sep", 10: "oct", 11: "nov", 12: "dec"}
    bad = 0
    for (cen, m), k in br.groupby([br.cen, br.c.dt.month]).size().items():
        if figs["ask_breaks.%s.%s" % (cen.lower(), mon[m])] != k:
            bad += 1
    ok("Item 201 cells", bad == 0 and len(br) == figs["ask_breaks.total"], "%d breaks" % len(br))

    pe = pd.read_excel(os.path.join(target, "mabry_teague_pay_estimates_fy2025_fy2026.xlsx"), dtype={"ITEM": str})
    ap = pd.read_csv(os.path.join(target, "ap_payments_contract_23-0388.csv"), dtype=str)
    tot = pe.groupby(["ESTIMATE_NO", "REVISION"])["AMOUNT"].sum().round(2)
    paid = {r.INVOICE_NO: float(r.GROSS_AMOUNT) for r in ap.itertuples()}
    gross = [round(float(x), 2) for x in ap.GROSS_AMOUNT]
    keep, used = set(), set()
    for gv in gross:
        m = [k for k, v in tot.items() if abs(v - gv) < 0.005]
        if len(m) == 1:
            keep.add(m[0])
            used.add(gv)
    done = {e for e, rv in keep}
    rest = [k for k in tot.index if k[0] not in done]
    for gv in gross:
        if gv in used:
            continue
        m = [(a, b) for a in rest for b in rest if a[0] < b[0] and abs(round(tot[a] + tot[b], 2) - gv) < 0.005]
        if len(m) == 1:
            keep.update(m[0])
    ok("every estimate has one paid revision", len({e for e, rv in keep}) == len(set(pe.ESTIMATE_NO)) == len(keep), len(keep))
    pe = pe[[(e, rv) in keep for e, rv in zip(pe.ESTIMATE_NO, pe.REVISION)]]
    grp_of = {"0410": "asphalt_pavement", "0420": "concrete_pavement", "0430": "sidewalk_and_driveway",
              "0431": "sidewalk_and_driveway", "0440": "sod"}
    pe = pe[pe.ITEM.isin(grp_of)].copy()
    pe["d"] = pd.to_datetime(pe.WORK_COMPLETED)
    pe = pe[(pe.d >= "2025-10-01") & (pe.d <= "2026-09-30")]
    pe["q"] = ((pe.d.dt.month - 10) % 12) // 3 + 1
    pe["g"] = pe.ITEM.map(grp_of)
    bad = 0
    for (gname, q), v in pe.groupby(["g", "q"])["QTY"].sum().items():
        if figs["ask_restoration.%s.q%d" % (gname, q)] != int(10 * round(v / 10.0)):
            bad += 1
    ok("Items 301 to 305 cells", bad == 0, "")

    from pypdf import PdfReader
    t = " ".join(p.extract_text() or "" for p in PdfReader(os.path.join(target, "itb_27-0114_service_repair_restoration_bid_schedule.pdf")).pages)
    t = re.sub(r"\s+", " ", t)
    ok("bid schedule reads back (positive control)", "A leak is one unit however many work orders" in t and len(t) > 1500, len(t))
    t2 = " ".join(p.extract_text() or "" for p in PdfReader(os.path.join(target, "cip_2027_water_main_replacement.pdf")).pages)
    ok("CIP reads back", "Services renewed" in t2, "")
    gold = os.path.join(os.path.dirname(os.path.abspath(target)), "golden")
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(gold, "service_leak_forecast_2027.xlsx"), data_only=True)
    vals = [c.value for row in wb["2027 forecast"].iter_rows() for c in row]
    ok("golden workbook carries the call", figs["call.service_leak_repairs_2027"] in vals, "")
    ok("golden chart present", os.path.getsize(os.path.join(gold, "service_leaks_by_month.png")) > 20000, "")
    print("verify: %d FAIL" % len(FAILS))
    return 1 if FAILS else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(HERE), "target")))
