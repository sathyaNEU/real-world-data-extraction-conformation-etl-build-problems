"""task173 documents: Workforce Planning's capacity standard, the outsourcer's Schedule B, the training centre's
2027 class calendar, the outsourcer's 2026 deadline-period invoices, the ACD's report field guide, the CRM export's
notes and the planning thread. Their text is kept here as data so checks.py and verify.py can quote and execute it."""
from __future__ import annotations

import datetime as dt
import os
import re

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import A4, LETTER                                        # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import cm                                                    # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import common as K          # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/"
for _n, _f in (("LibSans", "liberation/LiberationSans-Regular.ttf"), ("LibSans-Bold", "liberation/LiberationSans-Bold.ttf")):
    pdfmetrics.registerFont(TTFont(_n, FONT_DIR + _f))
INK = colors.HexColor("#1f2328")
GREY = colors.HexColor("#5b636b")
ACCENT = colors.HexColor("#1d4e5f")
RULE = colors.HexColor("#b4bcc4")
HEAD = colors.HexColor("#e4ecef")
HCS_ACCENT = colors.HexColor("#6a2c3e")

F_STANDARD = "expert_capacity_planning_standard.docx"
F_SCHEDULE_B = "hcs_msa_schedule_b_2026.pdf"
F_CALENDAR = "training_calendar_2027.pdf"
F_INVOICES = "hcs_invoices_2026_deadline_period.pdf"
F_GUIDE = "acd_report_field_guide.pdf"
F_CRM_NOTES = "crm_export_notes.txt"
F_THREAD = "re_2027_deadline_booking.eml"
DOMAIN = "ashdale.ca"


def addr(name, domain=DOMAIN):
    a = name.lower().split()
    return "%s.%s@%s" % (a[0], a[-1], domain)


def sentences(paras):
    return [s for p in paras for s in re.split(r"(?<=\.)\s+(?=[A-Z0-9(])", p)]


# ---------------------------------------------------------------------------------------------------- the standard
f27a, f27b = K.fortnight(K.SEASON_PLAN)
STANDARD = [
    ("Scope", [
        "This standard sets how Workforce Planning sizes the English expert queue of Ashdale Tax Live for each "
        "season's deadline fortnight: the load the in-house bench can carry and the agent hours booked with the "
        "outsourcer for the rest. The French queue is planned by the Montreal team and is not covered here. Dates and "
        "times in this standard are Eastern Time.",
    ]),
    ("The deadline fortnight", [
        "The deadline fortnight is the 14 days that end on the first Sunday after the filing deadline for most "
        "individuals, which is 30 April or the next business day when 30 April falls on a weekend. For the 2027 "
        "season it runs from %s to %s." % (K.long_date(f27a), K.long_date(f27b)),
    ]),
    ("Load", [
        "The fortnight's load is the expert contacts on the English-queue cases that open in the fortnight: each "
        "case's opening contact and every follow-up contact on that case, whenever the follow-up comes. Follow-ups "
        "that come during the fortnight on cases opened before it belong to the earlier period.",
        "The fortnight's new cases are the demand plan's daily forecast of new expert cases for the English queue.",
    ]),
    ("Handle time", [
        "Handle time is the talk time and after-contact work the CRM records for an interaction.",
    ]),
    ("Capacity", [
        "Each in-house expert on the English queue in the fortnight is scheduled for %d hours, %d hours a week over "
        "six days. Planned shrinkage is %d per cent (breaks, coaching, team meetings and system time) and planned "
        "occupancy is %d per cent, so each expert gives %d x %.2f x %.2f = %s handle hours across the fortnight." % (
            K.SCHEDULE_HOURS, K.SCHEDULE_HOURS / 2, round(K.SHRINKAGE * 100), round(K.OCCUPANCY * 100),
            K.SCHEDULE_HOURS, 1 - K.SHRINKAGE, K.OCCUPANCY, ("%.3f" % K.HANDLE_PER_EXPERT)),
        "An expert in a first season takes contacts from their class's go-live date.",
        "Each agent hour billed by the outsourcer gives %.2f handle hours." % K.HANDLE_PER_BILLED,
    ]),
    ("The bench", [
        "A season's bench is the English-queue experts who have confirmed the recall and the seats in the classes "
        "the training centre has scheduled to go live before the fortnight, each seat planned as filled.",
    ]),
    ("The booking", [
        "The in-house bench is planned to carry all the load its handle hours can take, and the outsourcer is booked "
        "for the rest of the fortnight's load: the billed agent hours whose handle hours carry those contacts at the "
        "outsourcer's handle time. Outsourcer hours are ordered in whole blocks of 1,000 agent hours, rounded up to "
        "the next block.",
    ]),
    ("Service level", [
        "Service level is reported for each day and each week of the fortnight and is not used to size it: the "
        "deadline days run above any service level the bench could be sized to.",
        "A customer's wait runs from the moment they join the expert queue to the start of their conversation with an "
        "expert, and the customer is counted on the date they joined the queue. Service level is the share of "
        "answered customers whose wait was three minutes or less.",
    ]),
    ("Review", [
        "Workforce Planning reviews this standard after each season. Version 3.2 replaces version 3.1 of "
        "9 October 2025.",
    ]),
]
STANDARD_SENTENCES = sentences([p for _, ps in STANDARD for p in ps])

# ---------------------------------------------------------------------------------------------------- Schedule B
SCHEDULE_B = [
    ("B.1 Rate", [
        "Deadline-period agent hours in the 2026 season are billed at CAD 38.50 an hour. HST is added at the New "
        "Brunswick rate.",
    ]),
    ("B.2 Billable time", [
        "Billable time is the time an HCS agent is signed in to the Client's expert queue in the Available, On Contact "
        "or After-Contact Work state, as the Client's ACD records it. Time in a Not Ready state is not billable, and "
        "team leads' time is included in the hourly rate.",
    ]),
    ("B.3 Invoices", [
        "HCS invoices weekly, for the week Monday to Sunday, within five business days of the week's end. Hours are "
        "taken from HCS's time-and-attendance records, net of meal and rest breaks, and each invoice lists the hours "
        "billed for each HCS agent by HCS agent ID.",
        "HCS's records and invoices are kept in Atlantic Time. The Client's ACD records time in UTC.",
    ]),
    ("B.4 Audit and credit", [
        "The Client may audit an invoice against its ACD records within 180 days of the invoice date. Where the hours "
        "an invoice bills for an agent in a week exceed that agent's billable time in the Client's ACD records for the "
        "same week, HCS credits the difference at the hourly rate.",
    ]),
    ("B.5 Deadline-period orders", [
        "The Client orders deadline-period agent hours for a season by 15 January, in whole blocks of 1,000 agent "
        "hours. An order cannot be reduced after 15 January. HCS staffs added blocks on two weeks' notice where it "
        "can.",
    ]),
    ("B.6 Access", [
        "The Client issues each HCS agent a login to its ACD before the agent's first shift. Logins are personal and "
        "are not shared or reassigned.",
    ]),
]
SCHEDULE_B_SENTENCES = sentences([p for _, ps in SCHEDULE_B for p in ps])

# ---------------------------------------------------------------------------------------------------- the calendar
CLASS_TABLE = [
    dict(code="27-1", room="Room A", train_start=dt.date(2027, 1, 18), golive=dt.date(2027, 2, 15), seats=60),
    dict(code="27-2", room="Room B", train_start=dt.date(2027, 2, 1), golive=dt.date(2027, 3, 1), seats=60),
    dict(code="27-3", room="Room A", train_start=dt.date(2027, 3, 1), golive=dt.date(2027, 3, 29), seats=60),
    dict(code="27-4", room="Room B", train_start=dt.date(2027, 3, 15), golive=dt.date(2027, 4, 12), seats=60),
]
CLASS_2026 = [("26-1", dt.date(2026, 1, 12), 60), ("26-2", dt.date(2026, 1, 26), 60), ("26-3", dt.date(2026, 3, 23), 40)]
CALENDAR_NOTES = [
    "Each class trains for four weeks, Monday to Friday, at the Waterloo training centre and goes live on the English "
    "queue on the Monday after its last training day.",
    "Seats are the class sizes recruiting is filling for. Offers for classes 27-1 and 27-2 go out from 9 November; "
    "classes 27-3 and 27-4 recruit from January.",
]


def _pdf_styles(accent=ACCENT):
    return dict(
        h1=ParagraphStyle("h1", fontName="LibSans-Bold", fontSize=15, leading=19, textColor=accent, spaceAfter=3),
        h2=ParagraphStyle("h2", fontName="LibSans-Bold", fontSize=10.5, leading=14, textColor=accent, spaceBefore=8,
                          spaceAfter=3),
        sub=ParagraphStyle("sub", fontName="LibSans", fontSize=8.8, leading=11.6, textColor=GREY, spaceAfter=10),
        body=ParagraphStyle("b", fontName="LibSans", fontSize=9.6, leading=13, textColor=INK, spaceAfter=6),
        cell=ParagraphStyle("c", fontName="LibSans", fontSize=8.4, leading=10.4, textColor=INK),
        cellr=ParagraphStyle("cr", fontName="LibSans", fontSize=8.4, leading=10.4, textColor=INK, alignment=2),
        cellb=ParagraphStyle("cb", fontName="LibSans-Bold", fontSize=8.4, leading=10.4, textColor=INK),
        cellbr=ParagraphStyle("cbr", fontName="LibSans-Bold", fontSize=8.4, leading=10.4, textColor=INK, alignment=2),
        small=ParagraphStyle("sm", fontName="LibSans", fontSize=7.9, leading=10.2, textColor=GREY, spaceAfter=4))


def _grid(data, widths, head_rows=1, accent_bg=HEAD):
    t = Table(data, colWidths=widths, repeatRows=head_rows)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, RULE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                           ("BACKGROUND", (0, 0), (-1, head_rows - 1), accent_bg), ("TOPPADDING", (0, 0), (-1, -1), 2.2),
                           ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2)]))
    return t


def _footer(text):
    def foot(canvas, doc):
        canvas.saveState()
        canvas.setFont("LibSans", 7.4)
        canvas.setFillColor(GREY)
        canvas.drawString(2.0 * cm, 1.2 * cm, text)
        canvas.drawRightString(doc.pagesize[0] - 2.0 * cm, 1.2 * cm, "Page %d" % doc.page)
        canvas.restoreState()
    return foot


def write_calendar(path):
    st = _pdf_styles()
    P = K.PEOPLE
    s = [Paragraph("Expert training calendar, 2027 season", st["h1"]),
         Paragraph("Ashdale Tax Live training centre, Waterloo. %s, training centre lead. Updated %s." % (
             P["training"], K.short_date(dt.date(2026, 10, 26))), st["sub"]),
         Paragraph("New-expert classes for the 2027 filing season", st["h2"])]
    head = ["Class", "Room", "Training starts", "Goes live", "Seats"]
    data = [[Paragraph(h, st["cellb"]) for h in head]]
    for c in CLASS_TABLE:
        data.append([Paragraph(c["code"], st["cell"]), Paragraph(c["room"], st["cell"]),
                     Paragraph(K.short_date(c["train_start"]), st["cell"]), Paragraph(K.short_date(c["golive"]), st["cell"]),
                     Paragraph(str(c["seats"]), st["cellr"])])
    s.append(_grid(data, [2.0 * cm, 2.2 * cm, 3.6 * cm, 3.6 * cm, 1.8 * cm]))
    s.append(Spacer(1, 8))
    for n in CALENDAR_NOTES:
        s.append(Paragraph(n, st["body"]))
    s.append(Paragraph("For reference: the 2026 classes", st["h2"]))
    data = [[Paragraph(h, st["cellb"]) for h in ("Class", "Went live", "Experts")]]
    for code, gl, n in CLASS_2026:
        data.append([Paragraph(code, st["cell"]), Paragraph(K.short_date(gl), st["cell"]), Paragraph(str(n), st["cellr"])])
    s.append(_grid(data, [2.0 * cm, 3.6 * cm, 1.8 * cm]))
    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=2.0 * cm, rightMargin=2.0 * cm, topMargin=1.8 * cm,
                            bottomMargin=1.8 * cm, title="Expert training calendar, 2027 season", author=P["training"],
                            subject="", creator="", producer="")
    doc.build(s, onFirstPage=_footer("Ashdale Software Inc.  |  Ashdale Tax Live training centre  |  2027 calendar"),
              onLaterPages=_footer("Ashdale Software Inc.  |  Ashdale Tax Live training centre  |  2027 calendar"))


def write_schedule_b(path):
    st = _pdf_styles(HCS_ACCENT)
    s = [Paragraph("Schedule B: Fees, invoicing and audit", st["h1"]),
         Paragraph("Master Services Agreement between Ashdale Software Inc. (the Client) and Hillsborough Contact "
                   "Services Ltd. (HCS), Schedule B as amended for the 2026 deadline period, 2 March 2026.", st["sub"])]
    for head, paras in SCHEDULE_B:
        s.append(Paragraph(head, st["h2"]))
        for p in paras:
            s.append(Paragraph(p, st["body"]))
    s.append(Spacer(1, 10))
    s.append(Paragraph("Signed for the Client: %s, Vice-President, Expert Services. Signed for HCS: %s, Account "
                       "Manager." % (K.PEOPLE["vp"], K.PEOPLE["vendor"]), st["small"]))
    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=2.0 * cm,
                            bottomMargin=2.0 * cm, title="MSA Schedule B 2026", author="Hillsborough Contact Services",
                            subject="", creator="", producer="")
    foot = _footer("Hillsborough Contact Services Ltd.  |  Ashdale Software Inc. MSA  |  Schedule B (2026)")
    doc.build(s, onFirstPage=foot, onLaterPages=foot)


def _money(x):
    return "{:,.2f}".format(x)


def write_invoices(path, lines):
    """The outsourcer's invoices for the three weeks of the 2026 deadline period, and the revised week-3 invoice."""
    import acd as A
    st = _pdf_styles(HCS_ACCENT)
    s = []
    items = [(no, wk, issued, None) for no, wk, issued in A.INVOICES] + [A.REVISED]
    for k, (no, wk, issued, replaces) in enumerate(items):
        if k:
            s.append(PageBreak())
        s.append(Paragraph("Hillsborough Contact Services Ltd.", st["h1"]))
        s.append(Paragraph("Contact centre services, Moncton, New Brunswick.", st["sub"]))
        s.append(Paragraph("INVOICE %s" % no, st["h2"]))
        meta = [["Invoice date", K.short_date(issued)], ["Bill to", "Ashdale Software Inc., Accounts Payable, "
                "Waterloo, Ontario"], ["Program", "Ashdale Tax Live, English expert queue, "
                "deadline period 2026"], ["Service week", "%s to %s (Atlantic Time)" % (
                    K.short_date(wk), K.short_date(wk + dt.timedelta(6)))], ["Terms", "Net 30"]]
        if replaces:
            meta.append(["Note", "This invoice replaces invoice %s. Hours are restated from time-and-attendance "
                                 "records." % replaces])
        s.append(_grid([[Paragraph(a, st["cellb"]), Paragraph(b, st["cell"])] for a, b in meta], [3.2 * cm, 13.0 * cm],
                       head_rows=0))
        s.append(Spacer(1, 6))
        tiny = ParagraphStyle("t", fontName="LibSans", fontSize=7.4, leading=8.6, textColor=INK)
        tinyr = ParagraphStyle("tr", fontName="LibSans", fontSize=7.4, leading=8.6, textColor=INK, alignment=2)
        tinyb = ParagraphStyle("tb", fontName="LibSans-Bold", fontSize=7.4, leading=8.6, textColor=INK)
        tinybr = ParagraphStyle("tbr", fontName="LibSans-Bold", fontSize=7.4, leading=8.6, textColor=INK, alignment=2)
        data = [[Paragraph(h, tinyb) for h in ("HCS agent ID", "Hours", "Rate", "Amount")]]
        tot_h = 0.0
        for emp, h in lines[no]:
            tot_h += h
            data.append([Paragraph(emp, tiny), Paragraph("%.2f" % h, tinyr),
                         Paragraph("%.2f" % A.RATE, tinyr), Paragraph(_money(h * A.RATE), tinyr)])
        sub = round(sum(round(h * A.RATE, 2) for _, h in lines[no]), 2)
        hst = round(sub * A.HST, 2)
        data.append([Paragraph("Total hours", tinyb), Paragraph("%.2f" % tot_h, tinybr), "",
                     Paragraph(_money(sub), tinybr)])
        data.append([Paragraph("HST 15%", tiny), "", "", Paragraph(_money(hst), tinyr)])
        data.append([Paragraph("Amount due (CAD)", tinyb), "", "", Paragraph(_money(sub + hst), tinybr)])
        t = Table(data, colWidths=[4.0 * cm, 2.6 * cm, 2.6 * cm, 3.6 * cm], repeatRows=1)
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, RULE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("BACKGROUND", (0, 0), (-1, 0), HEAD), ("TOPPADDING", (0, 0), (-1, -1), 0.6),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 0.9)]))
        s.append(t)
        s.append(Spacer(1, 4))
        s.append(Paragraph("Hours are billed from HCS time-and-attendance records, net of meal and rest breaks, per "
                           "Schedule B. Questions: %s, Account Manager, %s." % (
                               K.PEOPLE["vendor"], addr(K.PEOPLE["vendor"], "hillsboroughcs.ca")), st["small"]))
    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=1.3 * cm,
                            bottomMargin=1.6 * cm, title="HCS invoices, Ashdale deadline period 2026",
                            author="Hillsborough Contact Services", subject="", creator="", producer="")
    foot = _footer("Hillsborough Contact Services Ltd.  |  Accounts receivable")
    doc.build(s, onFirstPage=foot, onLaterPages=foot)


# ---------------------------------------------------------------------------------------------------- the ACD guide
GUIDE_STATES = [
    ("Available", "Signed in and ready for the queue to offer a contact."),
    ("On Contact", "From accepting an offer to the end of the conversation."),
    ("After-Contact Work", "After the conversation, until the agent goes ready again."),
    ("Not Ready", "Signed in but not taking offers. The reason code says why: Login, Break, Meal, Personal, No Answer, "
                  "Team Lead, End of Shift."),
    ("Logged Out", "The end of a session. The reason code says how the session ended: Agent logout, Session timeout or "
                   "Forced."),
]
GUIDE_STATE_COLS = [
    ("Agent Login", "The login the session was signed in with."),
    ("Agent Name", "The login's display name."),
    ("Session ID", "The ACD session, from sign-in to log-out."),
    ("State, Reason Code", "As in the table above."),
    ("Start (UTC), End (UTC)", "The part of the state the row covers."),
    ("Duration (s)", "The state's length, in seconds."),
]
GUIDE_CONTACT_COLS = [
    ("Contact ID", "The contact the queue routed."),
    ("Session ID", "The customer's session on the expert queue."),
    ("Queue, Channel", "The queue and the channel the customer chose."),
    ("Queued (UTC)", "When the contact was placed in the queue."),
    ("Offered (UTC)", "When the queue offered the contact to the agent on the row."),
    ("Agent Login", "The agent offered the contact."),
    ("Result", "Answered (the agent accepted the offer), No answer (the offer rang for 15 seconds without being "
               "accepted) or Abandoned (the customer left the queue)."),
    ("Answered (UTC)", "When the agent accepted the offer."),
    ("Ended (UTC)", "When the conversation ended, the offer stopped ringing or the customer left the queue."),
    ("CRM Interaction", "The CRM interaction opened from the contact, where there is one."),
]
GUIDE_NOTES = [
    "Both reports are exported in UTC. Ashdale's queues run on Eastern Time.",
    "An agent who does not accept an offer is set to Not Ready with the reason No Answer for a minute.",
]


def write_guide(path):
    st = _pdf_styles()
    P = K.PEOPLE
    s = [Paragraph("ACD reports: field guide", st["h1"]),
         Paragraph("Contact centre systems, Ashdale Software Inc. %s. Version 2.0, %s." % (
             P["crm"], K.short_date(dt.date(2026, 3, 9))), st["sub"]),
         Paragraph("Agent states", st["h2"])]
    data = [[Paragraph(h, st["cellb"]) for h in ("State", "Meaning")]]
    for a, b in GUIDE_STATES:
        data.append([Paragraph(a, st["cell"]), Paragraph(b, st["cell"])])
    s.append(_grid(data, [3.6 * cm, 12.6 * cm]))
    s.append(Paragraph("Agent state detail", st["h2"]))
    data = [[Paragraph(h, st["cellb"]) for h in ("Field", "Meaning")]]
    for a, b in GUIDE_STATE_COLS:
        data.append([Paragraph(a, st["cell"]), Paragraph(b, st["cell"])])
    s.append(_grid(data, [3.6 * cm, 12.6 * cm]))
    s.append(Paragraph("Contact detail", st["h2"]))
    data = [[Paragraph(h, st["cellb"]) for h in ("Field", "Meaning")]]
    for a, b in GUIDE_CONTACT_COLS:
        data.append([Paragraph(a, st["cell"]), Paragraph(b, st["cell"])])
    s.append(_grid(data, [3.6 * cm, 12.6 * cm]))
    s.append(Spacer(1, 6))
    for n in GUIDE_NOTES:
        s.append(Paragraph(n, st["body"]))
    doc = SimpleDocTemplate(path, pagesize=LETTER, leftMargin=2.0 * cm, rightMargin=2.0 * cm, topMargin=1.8 * cm,
                            bottomMargin=1.8 * cm, title="ACD reports: field guide", author=P["crm"], subject="",
                            creator="", producer="")
    foot = _footer("Ashdale Software Inc.  |  Contact centre systems  |  ACD field guide")
    doc.build(s, onFirstPage=foot, onLaterPages=foot)


GUIDE_SENTENCES = [b for _, b in GUIDE_STATES + GUIDE_STATE_COLS + GUIDE_CONTACT_COLS] + GUIDE_NOTES


# ---------------------------------------------------------------------------------------------------- CRM notes
def crm_notes(n_rows):
    P = K.PEOPLE
    L = [
        "Ashdale Tax Live CRM, saved export expert-interactions-2026",
        "Run 2026-10-29 21:40 Eastern by %s (CRM administrator), %s rows" % (P["crm"], "{:,}".format(n_rows)),
        "",
        "Scope: every expert interaction of the 2026 filing season, 12 January to 31 May 2026, both queues",
        "(Live Experts EN and Live Experts FR). Times are UTC.",
        "",
        "Interaction ID       one interaction between a customer and an expert",
        "Case Number          a case groups a customer's expert interactions on one return; the case's first",
        "                     interaction opens it",
        "Queue, Channel       the queue the customer came through and the channel used",
        "Interaction Type     New case for the interaction that opened the case, Follow-up for each later one",
        "Topic                the expert's topic for the interaction",
        "Started (UTC)        when the conversation began",
        "Ended (UTC)          when the expert closed the interaction, after the after-contact work",
        "Talk Seconds         the conversation",
        "After-Contact Work Seconds   notes and follow-up actions after the conversation",
        "Handled By           the expert's employee number, or the outsourcer agent's HCS agent ID",
        "Disposition          the expert's wrap-up code",
        "",
        "Employee numbers are the HR numbers used from 2026; seasonal_expert_history.csv carries the agent code",
        "each expert had on the old ticketing system.",
    ]
    return "\r\n".join(L) + "\r\n"


# ---------------------------------------------------------------------------------------------------- the thread
def _wrap(p, width=76):
    words, lines, cur = p.split(), [], ""
    for w in words:
        if cur and len(cur) + 1 + len(w) > width:
            lines.append(cur)
            cur = w
        else:
            cur = (cur + " " + w) if cur else w
    if cur:
        lines.append(cur)
    return lines


THREAD = [
    ("vp", "Mon, 26 Oct 2026 08:52:11 -0400", [
        "Kimberly,",
        "Please bring the Hillsborough booking for the 2027 deadline fortnight to the leadership meeting on Monday "
        "2 November. Tammy needs the order by 15 January and I would rather settle it before the classes start "
        "hiring.",
        "Mario",
    ]),
    ("training", "Mon, 26 Oct 2026 11:20:45 -0400", [
        "Kimberly, Mario,",
        "The 2027 calendar is updated on the shared drive: four classes of 60, going live on 15 February, 1 March, "
        "29 March and 12 April. Recruiting is filling all four and expects to, so please plan on the seats.",
        "Michelle",
    ]),
    ("wfm", "Thu, 29 Oct 2026 16:05:32 -0400", [
        "Mario, Michelle,",
        "The recall closed on Friday the 23rd and the roster now shows every answer: 150 of last season's English "
        "experts confirmed for 2027. I am pulling the 2026 CRM export, the ACD reports and Hillsborough's 2026 "
        "invoices today and will have the booking ready for Monday.",
        "Kimberly",
    ]),
]


def thread_text():
    import quopri
    P = K.PEOPLE
    body = []
    last = THREAD[-1]
    for p in last[2]:
        body += _wrap(p) + [""]
    body += ["-- ", P["wfm"], "Workforce Planning Manager, Ashdale Tax Live", ""]
    for who, date, paras in reversed(THREAD[:-1]):
        body.append("On %s, %s <%s> wrote:" % (date[:16], P[who], addr(P[who])))
        for p in paras:
            body += ["> " + ln for ln in _wrap(p, 74)] + [">"]
        body.append("")
    text = "\r\n".join(body).rstrip() + "\r\n"
    qp = quopri.encodestring(text.encode("utf-8")).decode("ascii")
    qp = qp.replace("\r\n", "\n").replace("\n", "\r\n")
    hdr = [
        "Return-Path: <%s>" % addr(P["wfm"]),
        "Message-ID: <MW4PR04MB7356A2F0C1D8E3B9A41C7@MW4PR04MB7356.namprd04.prod.outlook.com>",
        "In-Reply-To: <MW4PR04MB7356B91E25F04C6D8A2E9@MW4PR04MB7356.namprd04.prod.outlook.com>",
        "References: <YT3PR01MB8812C4A9E7F21B05D63A8@YT3PR01MB8812.CANPRD01.PROD.OUTLOOK.COM> "
        "<MW4PR04MB7356B91E25F04C6D8A2E9@MW4PR04MB7356.namprd04.prod.outlook.com>",
        "Date: %s" % last[1],
        "From: %s <%s>" % (P["wfm"], addr(P["wfm"])),
        "To: %s <%s>, %s <%s>" % (P["vp"], addr(P["vp"]), P["training"], addr(P["training"])),
        "Subject: RE: 2027 deadline booking",
        "MIME-Version: 1.0",
        "Content-Type: text/plain; charset=\"UTF-8\"",
        "Content-Transfer-Encoding: quoted-printable",
    ]
    return "\r\n".join(hdr) + "\r\n\r\n" + qp


THREAD_SENTENCES = sentences([p for _, _, ps in THREAD for p in ps])


# ---------------------------------------------------------------------------------------------------- the manifest
F_MANIFEST = "data_pull_manifest_2026-10-30.csv"
MANIFEST_COLS = ["File", "System", "Pulled by", "Pulled (Eastern)", "Covers", "Rows or pages"]


def manifest_rows(n, out):
    import acd as A
    import records as RC
    P = K.PEOPLE
    items = [
        (RC.F_CRM, "Ashdale Tax Live CRM", P["crm"], "2026-10-29 21:40", "2026 filing season, both expert queues",
         n["crm"]),
        (DOC_SELF + "crm_export_notes.txt", "Ashdale Tax Live CRM", P["crm"], "2026-10-29 21:44",
         "notes on the CRM export", None),
        (A.F_USERS, "ACD administration", P["crm"], "2026-10-29 16:12", "every ACD login", n["users"]),
        (A.F_STATES, "ACD reporting", P["crm"], "2026-10-29 16:20", "logins in the HCS Moncton team, 13 April to "
         "3 May 2026", n["states"]),
        (A.F_CONTACTS, "ACD reporting", P["crm"], "2026-10-29 16:31", "Live Experts EN, customers who joined the "
         "queue 20 April to 3 May 2026", n["contacts"]),
        (RC.F_ROSTER, "Workforce Planning roster", P["wfm"], "2026-10-29 16:31", "2026 seasonal experts and the 2027 "
         "recall to 23 October 2026", n["roster"]),
        (RC.F_HISTORY, "HR seasonal records", P["wfm"], "2026-10-27 10:05", "seasons 2021 to 2026", n["history"]),
        (RC.F_OLD_A, "old ticketing system archive", P["wfm"], "2026-10-27 10:22", "seasons 2021 and 2022",
         n["old_a"]),
        (RC.F_OLD_B, "old ticketing system archive", P["wfm"], "2026-10-27 10:24", "seasons 2023 to 2025",
         n["old_b"]),
        (RC.F_DEMAND, "Demand Planning", P["demand"], "2026-10-16 14:02", "2027 season daily forecast", n["demand"]),
        (A.F_ONBOARD, "Hillsborough Contact Services", P["vendor"], "2026-04-08 13:55", "agents on the 2026 "
         "deadline program", n["onboard"]),
        (F_INVOICES, "Accounts Payable", P["finance"], "2026-10-28 11:40", "HCS invoices for the 2026 deadline period",
         None),
        (F_SCHEDULE_B, "Contracts", P["finance"], "2026-10-28 11:42", "MSA Schedule B as amended 2 March 2026", None),
        ("csat_survey_responses_2026.csv", "customer survey platform", P["crm"], "2026-10-29 17:03", "responses to "
         "the 2026 post-contact survey", n["csat"]),
        ("finance_live_returns_weekly.csv", "Finance reporting", P["finance"], "2026-09-30 09:15", "weekly Live "
         "returns, 2021 to 2026 seasons", n["returns"]),
        (F_STANDARD, "Workforce Planning", P["wfm"], "2026-10-27 09:30", "version 3.2", None),
        (F_CALENDAR, "Training centre", P["training"], "2026-10-26 15:10", "2027 classes", None),
        (F_GUIDE, "Contact centre systems", P["crm"], "2026-10-29 16:35", "version 2.0", None),
        (F_THREAD, "email", P["wfm"], "2026-10-29 16:06", "2027 booking thread", None),
    ]
    rows = []
    for f, sysname, who, when, covers, rows_ in items:
        if rows_ is None:
            p = os.path.join(out, f)
            if f.endswith(".pdf"):
                from pypdf import PdfReader
                rows_ = "%d pages" % len(PdfReader(p).pages)
            elif f.endswith(".docx"):
                rows_ = "document"
            else:
                rows_ = "text"
        rows.append([f, sysname, who, when, covers, rows_])
    return rows


DOC_SELF = ""


# ---------------------------------------------------------------------------------------------------- Word
def _page_field(par):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        r = par.add_run()
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        elif text.strip() == "PAGE":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            r._r.append(it)
        else:
            r.text = text


def write_standard(path):
    import docx
    from docx.shared import Cm, Pt, RGBColor
    doc = docx.Document()
    stn = doc.styles["Normal"]
    stn.font.name = "Calibri"
    stn.font.size = Pt(10.5)
    stn.paragraph_format.space_after = Pt(6)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.left_margin = sec.right_margin = Cm(2.4)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    fp = sec.footer.paragraphs[0]
    fp.text = "Ashdale Tax Live, Workforce Planning, expert capacity planning standard v3.2, page "
    _page_field(fp)
    for r in fp.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x5B, 0x63, 0x6B)
    blue = RGBColor(0x1D, 0x4E, 0x5F)
    p = doc.add_paragraph()
    r = p.add_run("Expert capacity planning standard: English expert queue")
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = blue
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("Workforce Planning, Ashdale Tax Live. %s, version 3.2, 14 September 2026." % K.PEOPLE["wfm"])
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x5B, 0x63, 0x6B)
    for head, paras in STANDARD:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(head)
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = blue
        for tx in paras:
            doc.add_paragraph(tx)
    doc.save(path)


def write_all(out, n, invoice_lines):
    import postprocess as PP
    p = lambda f: os.path.join(out, f)
    P = K.PEOPLE
    with open(p(F_THREAD), "w", encoding="utf-8", newline="") as fh:
        fh.write(thread_text())
    with open(p(F_CRM_NOTES), "w", encoding="utf-8", newline="") as fh:
        fh.write(crm_notes(n["crm"]))
    write_standard(p(F_STANDARD))
    PP.normalize_docx(p(F_STANDARD), P["wfm"], P["wfm"], "2023-09-18T14:02:00Z", "2026-09-14T19:41:00Z",
                      revision=27, total_time=612, seed="task173/standard",
                      title="Expert capacity planning standard", company=K.COMPANY)
    write_calendar(p(F_CALENDAR))
    PP.normalize_pdf(p(F_CALENDAR), b"D:20261026150412-04'00'", b"D:20261026150412-04'00'")
    write_schedule_b(p(F_SCHEDULE_B))
    PP.normalize_pdf(p(F_SCHEDULE_B), b"D:20260302104455-04'00'", b"D:20260302104455-04'00'")
    write_invoices(p(F_INVOICES), invoice_lines)
    PP.normalize_pdf(p(F_INVOICES), b"D:20260519093120-03'00'", b"D:20260519093120-03'00'")
    write_guide(p(F_GUIDE))
    PP.normalize_pdf(p(F_GUIDE), b"D:20260309132207-04'00'", b"D:20260309132207-04'00'")
    K.write_csv(p(F_MANIFEST), MANIFEST_COLS, manifest_rows(n, out))
