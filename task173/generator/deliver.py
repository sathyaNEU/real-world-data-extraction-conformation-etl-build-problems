"""The golden deliverable (irrigation_forecast_2027.xlsx), generator/figures.json and submission.md, every figure
computed by calc.py from the shipped pack."""
from __future__ import annotations

import json
import os

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import calc
import common as C
import pack as PK

MONTHS = {4: "April", 5: "May", 6: "June", 7: "July", 8: "August", 9: "September"}
MON3 = {4: "Apr", 5: "May", 6: "Jun", 7: "Jul", 8: "Aug", 9: "Sep"}
DPNAME = {"RIVERSIDE": "Riverside", "MILL SLOUGH": "Mill Slough", "BENCH": "Bench", "CORWIN ROAD": "Corwin Road"}


def rnd(x, unit):
    return int(unit * round(x / unit))


def compute(target):
    P = calc.Pack(target)
    g = calc.golden(P)
    A = calc.ask_a(P)
    Cc = calc.ask_c(P)
    n, e = calc.ask_b(P)
    det = g["detail"]
    fig = {
        "call": {"forecast_mwh": rnd(g["total_kwh"] / 1000.0, C.CALL_UNIT)},
        "comp_bench_mwh": rnd(g["juniors_kwh"] / 1000.0, 1000),
        "comp_bottomland_well_af": rnd(float(det.loc[det["well"], "cut"].sum()), 5000),
        "normal_eto_in": round(g["etn"], 2),
        "order_limit_pct": int(round(100 * C.SENIOR_LIMIT_2027)),
        "call_rounding_mwh": C.CALL_UNIT,
        "dst_day_intervals": int((P.dpi[P.dpi["interval_end"].str.startswith("2026-03-08")].groupby("meter_no").size()).min()),
    }
    for dp in C.DPS:
        key = dp.replace(" ", "_")
        fig["a_" + key] = {"2026-%02d" % m: int(round(A[(dp, m)] * 1000)) for m in C.SEASON}
    for dp in C.DPS:
        key = dp.replace(" ", "_")
        fig["b_" + key] = {"2026-%02d" % m: int(round(Cc[(dp, m)])) for m in C.SEASON}
    for dp in C.DPS:
        key = dp.replace(" ", "_")
        fig["c_" + key] = {"services": int(n[dp]), "billed_mwh": round(e[dp], 1)}
    return P, g, A, Cc, n, e, fig


THIN = Side(style="thin", color="9AA4AE")
HEAD_FILL = PatternFill("solid", fgColor="DCE3EA")


def _hdr(ws, row, cols):
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)


def write_xlsx(path, P, g, A, Cc, n, e, fig):
    wb = Workbook()
    ws = wb.active
    ws.title = "2027 irrigation forecast"
    ws["A1"] = "Varner Valley Public Utility District: Schedule IP (irrigation pumping) energy, 1 April to 30 September 2027"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Halvern form MLF-3. Normal season (station 247, mean April to September reference ET 2007 to 2026, "
                "%.2f in); surface supply under Watermaster Order No. 2026-14; services in service on 30 September 2026." % g["etn"])
    ws["A4"] = "Forecast, MWh at the retail meter (nearest 5,000)"
    ws["B4"] = fig["call"]["forecast_mwh"]
    ws["B4"].number_format = "#,##0"
    ws["A4"].font = Font(bold=True)
    ws["B4"].font = Font(bold=True)
    det = g["detail"]
    rows = [
        ("Lone Bench services (bench wells), 2022 season at normal ET, MWh", g["juniors_kwh"] / 1000.0, "#,##0"),
        ("Bottomland plants, normal-season energy, MWh", g["seniors_normal_kwh"] / 1000.0, "#,##0"),
        ("Bottomland diversion above a quarter of the allotment, acre-feet", float(det["cut"].sum()), "#,##0"),
        ("  of which made up from the plants' own wells, acre-feet", float(det.loc[det["well"], "cut"].sum()), "#,##0"),
        ("  of which left unirrigated (plants with no well), acre-feet", float(det.loc[~det["well"], "cut"].sum()), "#,##0"),
        ("Bottomland wells' added energy, MWh", float((det["cut"] * det["kw"] * det["well"]).sum()) / 1000.0, "#,##0"),
        ("Bottomland river lifts' energy no longer drawn, MWh", -float((det["cut"] * det["kr"]).sum()) / 1000.0, "#,##0"),
        ("Total before rounding, MWh", g["total_kwh"] / 1000.0, "#,##0"),
    ]
    ws["A6"] = "Working"
    ws["A6"].font = Font(bold=True)
    for i, (lab, v, f) in enumerate(rows, start=7):
        ws.cell(row=i, column=1, value=lab)
        c = ws.cell(row=i, column=2, value=int(round(v)))
        c.number_format = f
    r0 = 7 + len(rows) + 1
    ws.cell(row=r0, column=1, value="Services in the forecast: %d Lone Bench, %d bottomland (%d with a well at the plant)"
            % (g["n_juniors"], g["n_seniors"], g["n_wells"]))
    ws.cell(row=r0 + 1, column=1, value=("Bottomland energy per acre-foot at each plant's river lift and well: each plant's "
                                         "water-year energy fitted to its own diversion statements and its metered well "
                                         "extraction, water years 2017 to 2026."))
    ws.column_dimensions["A"].width = 78
    ws.column_dimensions["B"].width = 14
    ws.freeze_panes = "A4"

    # plant detail, so the working can be followed
    wd = wb.create_sheet("Bottomland plants")
    cols = ["Service", "Decree right", "Season allotment (AF)", "Normal-season diversion (AF)", "Diversion above the limit (AF)",
            "Well at the plant", "River lift kWh/AF", "Well kWh/AF", "Normal-season energy (MWh)", "2027 energy (MWh)"]
    _hdr(wd, 1, cols)
    for i, r in enumerate(det.sort_values("svc").itertuples(index=False), start=2):
        e27 = r.en + (r.cut * r.kw if r.well else 0.0) - r.cut * r.kr
        vals = [r.svc, r.right, int(r.allot), round(r.dn, 1), round(r.cut, 1), "Yes" if r.well else "No",
                round(r.kr, 1), round(r.kw, 1) if r.well else None, round(r.en / 1000.0, 1), round(e27 / 1000.0, 1)]
        for j, v in enumerate(vals, start=1):
            wd.cell(row=i, column=j, value=v)
    widths = [10, 12, 12, 14, 14, 10, 11, 11, 13, 12]
    for j, w in enumerate(widths, start=1):
        wd.column_dimensions[get_column_letter(j)].width = w
    wd.freeze_panes = "B2"

    wa = wb.create_sheet("MLF-4(a) DP energy 2026")
    wa["A1"] = "Energy Halvern metered at each delivery point, billing months April to September 2026 (kWh, Rate W-2 billing months)"
    wa["A1"].font = Font(bold=True)
    _hdr(wa, 3, ["Delivery point"] + [MONTHS[m] for m in C.SEASON])
    for i, dp in enumerate(C.DPS, start=4):
        wa.cell(row=i, column=1, value="%s (%s)" % (DPNAME[dp], C.DP_CODES[dp]))
        for j, m in enumerate(C.SEASON, start=2):
            c = wa.cell(row=i, column=j, value=fig["a_" + dp.replace(" ", "_")]["2026-%02d" % m])
            c.number_format = "#,##0"
    wa.cell(row=9, column=1, value=("Revenue meters 81-2207 (Riverside), 81-3318 (Mill Slough), 81-1940 and from 15 June 81-1977 "
                                    "(Bench, banks 1 and 2), 81-4062 (Corwin Road); the check meters 81-2251, 81-3362, 81-1988 and "
                                    "81-4107 not counted. Meter clocks run on Pacific standard time; intervals moved to prevailing "
                                    "time before the months were cut."))
    wa.column_dimensions["A"].width = 24
    for j in range(2, 8):
        wa.column_dimensions[get_column_letter(j)].width = 13

    wc = wb.create_sheet("MLF-4(b) peak demand 2026")
    wc["A1"] = "Demand at each delivery point in Halvern's monthly system peak hour, April to September 2026 (kW)"
    wc["A1"].font = Font(bold=True)
    _hdr(wc, 3, ["Delivery point"] + ["%s %s HE %s" % (MON3[int(r['month'][5:])], r["peak_date"][8:], r["hour_ending"][:2])
                                       for _, r in P.peaks.iterrows() if int(r["month"][5:]) in C.SEASON])
    for i, dp in enumerate(C.DPS, start=4):
        wc.cell(row=i, column=1, value="%s (%s)" % (DPNAME[dp], C.DP_CODES[dp]))
        for j, m in enumerate(C.SEASON, start=2):
            c = wc.cell(row=i, column=j, value=fig["b_" + dp.replace(" ", "_")]["2026-%02d" % m])
            c.number_format = "#,##0"
    wc.cell(row=9, column=1, value="Hour ending in Pacific prevailing time, as posted; the four 15-minute intervals of the clock hour summed.")
    wc.column_dimensions["A"].width = 24
    for j in range(2, 8):
        wc.column_dimensions[get_column_letter(j)].width = 15

    wb_ = wb.create_sheet("MLF-4(c) services by DP")
    wb_["A1"] = "Schedule IP services each delivery point supplied through the 2026 season, and their billed energy April to September 2026"
    wb_["A1"].font = Font(bold=True)
    _hdr(wb_, 3, ["Delivery point", "IP services", "Billed energy (MWh)"])
    for i, dp in enumerate(C.DPS, start=4):
        k = "c_" + dp.replace(" ", "_")
        wb_.cell(row=i, column=1, value="%s (%s)" % (DPNAME[dp], C.DP_CODES[dp]))
        wb_.cell(row=i, column=2, value=fig[k]["services"])
        c = wb_.cell(row=i, column=3, value=fig[k]["billed_mwh"])
        c.number_format = "#,##0.0"
    wb_.cell(row=9, column=1, value=("Each service at the delivery point its feeder was supplied from in the season (switching orders "
                                     "SO-25-207, feeder 3204 to Corwin Road from 3 November 2025; SO-26-031, feeder 1204 to Mill Slough "
                                     "from 16 March 2026)."))
    wb_.column_dimensions["A"].width = 24
    wb_.column_dimensions["B"].width = 12
    wb_.column_dimensions["C"].width = 20

    wn = wb.create_sheet("Notes")
    notes = [
        "Prepared for Halvern Power Authority form MLF-3 and MLF-4, 2027 planning year.",
        "Class: Schedule IP services in service on 30 September 2026 (%d services)." % (g["n_juniors"] + g["n_seniors"]),
        "Season: 1 April to 30 September 2027. Energy at the retail meter.",
        "Normal season: station 247, mean April to September reference ET of the 20 seasons 2007 to 2026 (%.2f in)." % fig["normal_eto_in"],
        "Surface supply: Varner River Watermaster Order No. 2026-14 (Lone Bench Irrigation District right suspended; riparian and "
        "pre-1914 diversions held to %d per cent of each right's season allotment)." % fig["order_limit_pct"],
        "Bills: IP billing extract of 16 October 2026; legacy accounts before the January 2020 conversion mapped through the crosswalk.",
        "Wells: Varner Subbasin Groundwater Agency extraction reports, water years 2017 to 2026 (water year 2026 as received by 15 October).",
        "Base-year tables: Halvern interval export and posted peaks for 2026; feeder sources from the OMS switching orders.",
    ]
    for line in notes:
        wn.append([line])
    wn.column_dimensions["A"].width = 120
    PK.save_xlsx(wb, path, "Varner Valley Public Utility District")


def fmt(v):
    return "{:,}".format(v)


def write_submission(path, P, g, A, Cc, n, e, fig):
    det = g["detail"]
    call = fig["call"]["forecast_mwh"]
    L = []
    L.append("## Tags\n")
    L.append("**Domain:** Energy Systems & Utilities (load forecasting).")
    L.append("**Analytical objective:** Forecasting & Predictive Modeling.\n")
    L.append("## 1. Final Recommendation\n")
    L.append("**File %s MWh as the irrigation class's April to September 2027 energy.** The Lone Bench services draw what "
             "they drew in 2022, when the district delivered nothing, and the bottomland plants make the diversion above a "
             "quarter of their allotments up from their own shallow wells, whose energy per acre-foot each plant's meter "
             "separates from its river lift's. Not the bench wells' energy per acre-foot carried onto the bottomland cut, "
             "which prices shallow water at deep lifts. Not each plant's energy over its water, which reads a river lift and a "
             "well on one meter as one pump.\n" % fmt(call))
    L.append("## 2. Critical Components\n")
    L.append("1. The Lone Bench services draw **%s MWh** in a normal season with no district water." % fmt(fig["comp_bench_mwh"]))
    L.append("2. The bottomland plants' wells make up **%s acre-feet** of diversion the order removes." % fmt(fig["comp_bottomland_well_af"]))
    L.append("")
    L.append("## 3. Step-by-Step Solution\n")
    L.append("1. Summed every read line in vvpud_ip_billing_2015-2026.csv by service and bill month, legacy accounts mapped "
             "through vvpud_cis_account_crosswalk_2020.csv, and kept the services ACTIVE in vvpud_ip_service_register_2026-09-30.xlsx; "
             "a service is a bottomland plant when varner_watermaster_diversion_statements_2015-2026.csv names it as a point's measure.")
    L.append("2. Took the April to September total of station247_daily_eto_2007-2026.csv in each season and their 2007 to 2026 "
             "mean, %.2f in, as the normal season the forecast instructions define." % fig["normal_eto_in"])
    L.append("3. Put the Lone Bench services at their 2022 April to September energy times the normal over 2022's ET: %s MWh."
             % fmt(fig["comp_bench_mwh"]))
    L.append("4. Put each bottomland plant at its 2015 to 2026 seasons' energy, each at normal ET, averaged, and its diversion "
             "the same way; the order holds diversion to %d per cent of the allotment in varner_decree_allotment_schedule_2026.xlsx, "
             "so the diversion above that is lost to the river and made up from a well where the plant has one in "
             "vsga_extraction_reports_wy2017-wy2026.xlsx: %s acre-feet." % (fig["order_limit_pct"], fmt(fig["comp_bottomland_well_af"])))
    L.append("5. Fitted each plant's October to September energy, without intercept, on its season diversion and its metered "
             "well extraction for water years 2017 to 2026 to get its river lift's and its well's kWh per acre-foot (a plant "
             "with no well: its energy over its diversion).")
    L.append("6. Added each plant's well energy on the water made up and took off its river lift's energy on the diversion "
             "lost; with the Lone Bench services the class draws %s MWh to the nearest %s." % (fmt(call), fmt(fig["call_rounding_mwh"])))
    L.append("7. Read hpa_dp_interval_export_2026.csv as meter standard time (every meter carries %d intervals on 8 March 2026), "
             "moved daylight-time intervals an hour forward, kept the revenue meters (each point's check meter repeats its revenue meter "
             "one interval early; Bench's second bank adds its own meter from 15 June) and summed each billing month, and each "
             "posted peak's clock hour from hpa_2026_monthly_system_peaks.csv." % fig["dst_day_intervals"])
    L.append("8. Put each service at the delivery point its feeder was supplied from through the season in "
             "vvpud_oms_feeder_sources.csv (two feeders moved by switching order, one in November 2025 and one in March 2026), "
             "not the GIS substation, and summed its April to September 2026 bills.\n")
    L.append("## 4. Deliverable Answers\n")
    L.append("### %s\n" % C.DELIVERABLE)
    L.append("1. Schedule IP energy, 1 April to 30 September 2027, normal season: %s MWh." % fmt(call))
    L.append("2. Energy Halvern metered at each delivery point, billing months April to September 2026, kWh:")
    for dp in C.DPS:
        a = fig["a_" + dp.replace(" ", "_")]
        L.append("   - %s: " % DPNAME[dp] + ", ".join("%s %s" % (MONTHS[m], fmt(a["2026-%02d" % m])) for m in C.SEASON))
    L.append("3. Demand at each delivery point in Halvern's monthly system peak hour, April to September 2026, kW:")
    for dp in C.DPS:
        c = fig["b_" + dp.replace(" ", "_")]
        L.append("   - %s: " % DPNAME[dp] + ", ".join("%s %s" % (MONTHS[m], fmt(c["2026-%02d" % m])) for m in C.SEASON))
    L.append("4. Irrigation services each delivery point supplied through the 2026 season and their April to September 2026 billed energy:")
    for dp in C.DPS:
        b = fig["c_" + dp.replace(" ", "_")]
        L.append("   - %s: %d services, %s MWh" % (DPNAME[dp], b["services"], "{:,.1f}".format(b["billed_mwh"])))
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(L) + "\n")


def write_figures(path, fig, g):
    out = dict(fig)
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        json.dump(out, f, indent=1, sort_keys=False)
        f.write("\n")
