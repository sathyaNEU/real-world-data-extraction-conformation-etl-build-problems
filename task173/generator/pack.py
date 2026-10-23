"""Write the evidence pack into a target folder: CSV, XLSX and PDF exports in the organisations' own layouts."""
from __future__ import annotations

import datetime as dt
import io
import os
import re
import zipfile

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

import common as C
import docs as D
import dpi
import records as R

STAMP = dt.datetime(2026, 10, 16, 15, 20, 0)
ZIP_DT = (2026, 10, 16, 15, 20, 0)


def write_csv(df, path):
    df.to_csv(path, index=False, lineterminator="\n")
    return len(df)


def save_xlsx(wb, path, creator):
    wb.properties.creator = creator
    wb.properties.lastModifiedBy = creator
    wb.properties.created = STAMP
    wb.properties.modified = STAMP
    buf = io.BytesIO()
    wb.save(buf)
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/app.xml":
                data = re.sub(rb"<Application>[^<]*</Application>", b"<Application>Microsoft Excel</Application>", data)
                data = re.sub(rb"<AppVersion>[^<]*</AppVersion>", b"<AppVersion>16.0300</AppVersion>", data)
            if item.filename == "docProps/core.xml":
                # openpyxl stamps the save clock as modified; the file carries the extract's own time
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                              rb"\g<1>" + STAMP.strftime("%Y-%m-%dT%H:%M:%SZ").encode() + rb"\g<2>", data)
            zi = zipfile.ZipInfo(item.filename, date_time=ZIP_DT)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            zout.writestr(zi, data)


def sheet(wb, title, df, widths=None, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and np.isnan(v)) else (v.item() if hasattr(v, "item") else v) for v in row])
    for j, col in enumerate(df.columns, start=1):
        w = (widths or {}).get(col) or min(max(len(str(col)), int(df[col].astype(str).str.len().max() if len(df) else 8)) + 2, 48)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    return ws


def write_all(W, ids, out):
    os.makedirs(out, exist_ok=True)
    n = {}
    bill = R.billing(W, ids)
    n[C.F_BILL] = write_csv(bill, os.path.join(out, C.F_BILL))

    reg = R.register(W, ids)
    wb = Workbook()
    sheet(wb, "IP services", reg, first=True)
    ws = wb.create_sheet("About")
    for line in ["Varner Valley Public Utility District, customer information system",
                 "Schedule IP (irrigation pumping) service register as of 30 September 2026",
                 "One row per service. A service closed before the register date stays on the register with its close date.",
                 "plant: the pumping plant the service supplies, as described at connection or at the last field check.",
                 "Run 16 October 2026 by Customer Accounting."]:
        ws.append([line])
    save_xlsx(wb, os.path.join(out, C.F_REG), "Varner Valley Public Utility District")
    n[C.F_REG] = len(reg)

    n[C.F_XW] = write_csv(R.crosswalk(ids), os.path.join(out, C.F_XW))
    n[C.F_SPT] = write_csv(R.service_points(W, ids), os.path.join(out, C.F_SPT))
    n[C.F_FDR] = write_csv(R.feeder_sources(ids), os.path.join(out, C.F_FDR))

    gsa = R.gsa_reports(W, ids)
    wb = Workbook()
    first = True
    for wy, df in gsa.items():
        if wy <= 2020:
            d2 = pd.DataFrame({"State Well No.": df["well"], "APN": df["apn"], "Owner of Record": df["owner"],
                               "Power Account": df["power"], "Volume (ac-ft)": df["af"].round(1),
                               "Measurement Method": df["method"]})
        else:
            d2 = pd.DataFrame({"water_year": wy, "state_well_no": df["well"], "apn": df["apn"], "owner": df["owner"],
                               "power_source": df["power"], "extraction_af": df["af"].round(2),
                               "method": df["method"]})
        d2 = d2.sort_values(d2.columns[1] if wy > 2020 else d2.columns[0]).reset_index(drop=True)
        sheet(wb, "WY%d" % wy, d2, first=first)
        first = False
    ws = wb.create_sheet("Notes")
    for line in ["Varner Subbasin Groundwater Agency: annual extraction reports, water years 2017 to 2026",
                 "Water year: 1 October to 30 September, named for the year it ends.",
                 "Every well extracting more than 2 acre-feet a year reports metered extraction (Agency Rule 3.2, from water year 2017).",
                 "Reports from water year 2021 are filed through the Agency's portal; earlier years were keyed from paper forms.",
                 "Water year 2026 reports were due 31 October 2026; this export holds the reports received by 15 October 2026."]:
        ws.append([line])
    save_xlsx(wb, os.path.join(out, C.F_GSA), "Varner Subbasin Groundwater Agency")
    n[C.F_GSA] = sum(len(v) for v in gsa.values())

    n[C.F_DIV] = write_csv(R.diversions(W, ids), os.path.join(out, C.F_DIV))

    al = R.allotments(W, ids)
    wb = Workbook()
    sheet(wb, "Schedule of allotments", al, first=True)
    ws = wb.create_sheet("Notes")
    for line in ["Varner River Watermaster: schedule of allotments under the Decree, as maintained for the 2026 season.",
                 "season_allotment_af: the right's allotment for an irrigation season (1 April to 30 September), acre-feet at the point of diversion.",
                 "place_of_use: assessor parcels on the right's place of use."]:
        ws.append([line])
    save_xlsx(wb, os.path.join(out, C.F_ALT), "Varner River Watermaster")
    n[C.F_ALT] = len(al)

    con, alc = R.lbid(W, ids)
    wb = Workbook()
    sheet(wb, "Contracts", con, first=True)
    sheet(wb, "Allocations", alc)
    save_xlsx(wb, os.path.join(out, C.F_LBID), "Lone Bench Irrigation District")
    n[C.F_LBID] = len(con)

    n[C.F_ETO] = write_csv(R.eto(W), os.path.join(out, C.F_ETO))

    d = dpi.build(W, ids)
    n[C.F_DPI] = write_csv(d, os.path.join(out, C.F_DPI))
    n[C.F_PEAK] = write_csv(dpi.peaks(W), os.path.join(out, C.F_PEAK))

    old, new = R.pump_tests(W, ids)
    wb = Workbook()
    sheet(wb, "2019-2021 Ostrom Pump Testing", old, first=True)
    sheet(wb, "2022-2026 Aldana Pump & Well", new)
    save_xlsx(wb, os.path.join(out, C.F_PTEST), "Varner Valley Public Utility District")
    n[C.F_PTEST] = len(old) + len(new)

    D.order(os.path.join(out, C.F_ORDER))
    D.instructions(os.path.join(out, C.F_INSTR))
    D.rate_excerpt(os.path.join(out, C.F_RATE))
    D.crop_report(os.path.join(out, C.F_CROP), crop_rows(W))
    write_dictionary(os.path.join(out, C.F_DICT))
    write_manifest(os.path.join(out, C.F_MAN), n)
    return n


def crop_rows(W):
    r = R.rng("crop")
    acres = {}
    for c, a in zip(W.J["crop"], W.J["acres"]):
        acres[c] = acres.get(c, 0) + a
    for c, a in zip(W.S["crop"], W.S["acres"]):
        acres[c] = acres.get(c, 0) + a
    yields = {"ALMOND": 1.12, "PISTACHIO": 1.31, "WALNUT": 1.86, "WINEGRAPE": 7.4, "CORN SILAGE": 27.5, "ALFALFA": 7.1,
              "IRRIGATED PASTURE": None, "PROCESSING TOMATO": 49.0, "PRUNE": 1.95}
    price = {"ALMOND": 4180, "PISTACHIO": 5390, "WALNUT": 2210, "WINEGRAPE": 560, "CORN SILAGE": 61, "ALFALFA": 248,
             "IRRIGATED PASTURE": None, "PROCESSING TOMATO": 104, "PRUNE": 2440}
    rows = []
    for c in sorted(acres):
        a = int(round(acres[c] * float(r.uniform(1.9, 2.6)), -1))   # the whole county, not only the PUD's territory
        y = yields[c]
        v = "" if y is None else "{:,}".format(int(round(a * y * price[c] * float(r.normal(1, 0.04)), -3)))
        rows.append([c.title(), "{:,}".format(a), "" if y is None else "%.2f" % y, v])
    return rows


DICT = """VARNER VALLEY PUBLIC UTILITY DISTRICT
Customer Accounting and Engineering: data dictionary for the October 2026 irrigation extract

vvpud_ip_billing_2015-2026.csv
  One row per meter read line on a Schedule IP bill, January 2015 to September 2026 bill months.
  acct_no       account number; bills before the January 2020 conversion carry the legacy number only
  svc_no        service number (blank on legacy rows; see the conversion crosswalk)
  bill_month    the month the read closes (YYYY-MM); IP meters are read at midnight at month end
  read_start, read_end   the read period
  meter_no      meter on the read line; a bill with a meter exchange has a line for each meter
  prev_read, curr_read, mult   register reads and the meter multiplier; kwh = (curr_read - prev_read) x mult
  read_type     A actual, E estimated (the next actual read carries the difference)
  bill_status   BILLED; CANCELLED (a bill later cancelled), CANCEL (its reversal), REBILL (the replacement)

vvpud_ip_service_register_2026-09-30.xlsx   one row per Schedule IP service; status and close date as of 30 September 2026
vvpud_cis_account_crosswalk_2020.csv        legacy account to new account and service number, January 2020 conversion
vvpud_gis_service_points_2026-09-30.csv     each IP service point's transformer, feeder and substation in the GIS
vvpud_oms_feeder_sources.csv                each feeder's source bus and delivery point, with effective dates, from switching orders

vsga_extraction_reports_wy2017-wy2026.xlsx   Varner Subbasin Groundwater Agency annual well extraction reports, one sheet per water year
varner_watermaster_diversion_statements_2015-2026.csv   monthly diversion at each point of diversion, acre-feet
  measure       how the diversion is measured (the electric service of a pumped point, or the canal gauge)
varner_decree_allotment_schedule_2026.xlsx   rights under the Varner River Decree and their season allotments
lbid_contracts_and_allocations_2015-2026.xlsx   Lone Bench Irrigation District landowner contracts and each season's allocation
station247_daily_eto_2007-2026.csv           daily reference evapotranspiration, inches, station 247 (Varner Valley); qc M marks a day filled from neighbouring stations

hpa_dp_interval_export_2026.csv   Halvern delivery-point interval data, 1 January to 15 October 2026, from the member portal
  dp_code       Halvern delivery point code (VV-01 Riverside, VV-02 Mill Slough, VV-03 Bench, VV-04 Corwin Road)
  meter_no      Halvern meter number
  interval_end  the 15-minute interval as the meter records it
  kwh, kvarh    energy delivered in the interval
  status        E where Halvern estimated the interval
hpa_2026_monthly_system_peaks.csv   Halvern's posted monthly system peaks, January to September 2026

vvpud_pump_efficiency_program_2019-2026.xlsx   pump efficiency tests under the PUD's test programme (plants of 75 hp and more)
"""


def write_dictionary(path):
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(DICT)


def write_manifest(path, n):
    rows = [
        (C.F_BILL, "VVPUD CIS (legacy billing to Jan 2020)", "2026-10-16", "Jan 2015 to Sep 2026 bill months", n[C.F_BILL]),
        (C.F_REG, "VVPUD CIS", "2026-10-16", "IP services as of 2026-09-30", n[C.F_REG]),
        (C.F_XW, "VVPUD CIS conversion", "2020-01-06", "accounts converted Jan 2020", ""),
        (C.F_SPT, "VVPUD GIS", "2026-10-16", "IP service points as of 2026-09-30", n[C.F_SPT]),
        (C.F_FDR, "VVPUD OMS", "2026-10-16", "feeder sources to date", n[C.F_FDR]),
        (C.F_GSA, "VSGA portal and paper reports", "2026-10-15", "WY2017 to WY2026 reports received", n[C.F_GSA]),
        (C.F_DIV, "Varner River Watermaster", "2026-10-14", "seasons 2015 to 2026", n[C.F_DIV]),
        (C.F_ALT, "Varner River Watermaster", "2026-10-14", "schedule as maintained for 2026", n[C.F_ALT]),
        (C.F_LBID, "Lone Bench Irrigation District", "2026-10-09", "contracts and allocations 2015 to 2026", n[C.F_LBID]),
        (C.F_ETO, "station 247", "2026-10-16", "2007-01-01 to 2026-10-15", n[C.F_ETO]),
        (C.F_DPI, "Halvern member portal", "2026-10-16", "2026-01-01 to 2026-10-15", n[C.F_DPI]),
        (C.F_PEAK, "Halvern member portal", "2026-10-16", "Jan to Sep 2026", n[C.F_PEAK]),
        (C.F_PTEST, "VVPUD energy services", "2026-10-16", "tests 2019 to 2026", n[C.F_PTEST]),
        (C.F_ORDER, "Varner River Watermaster", "2026-09-28", "Order No. 2026-14", ""),
        (C.F_INSTR, "Halvern", "2026-08-04", "MLF instructions 2027", ""),
        (C.F_RATE, "Halvern", "2025-01-01", "Rate Schedule W-2 excerpt", ""),
        (C.F_CROP, "Kestrel County", "2026-06-30", "2025 report excerpt", ""),
        (C.F_DICT, "VVPUD", "2026-10-16", "this extract", ""),
    ]
    df = pd.DataFrame(rows, columns=["file", "source", "extracted", "covers", "rows"])
    write_csv(df, path)
