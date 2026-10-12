"""The two deliverables: the per-filer reference matches (CSV) and the committee briefing deck (PPTX) with the
forecast-to-actual bridge. Figures come from calc.py; verify_pack.py checks them on its own path."""
import csv
import datetime as dt
import io
import re
import zipfile
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

import calc
import world as W

TASK = Path(__file__).resolve().parent.parent
T = TASK / "target"
G = TASK / "golden"
DECK = "fy2026_lodging_tax_variance.pptx"
CSVF = "fy2026_lodging_reference_matches.csv"
NAMES = {"CC": "Convention center closure", "LX": "Long-stay exemption", "CO": "County option cut (Prairie, Bluff)",
         "PC": "Facilitator collection (short-term rentals)", "AH": "Airline hub reduction"}
NAVY, GREY, RED, GREEN = "#2F4858", "#8A8F98", "#B5442B", "#3D7A5A"


def k(x):
    return int(round(x / 1000.0))


def figures():
    f, d = calc.all_figures(T)
    f["cash"] = calc.cash_reconciliation(d["pn"], d["ledger"])
    return f, d


# ---------------------------------------------------------------------------------------------- the CSV
def write_csv(f, d):
    pn = d["pn"]
    rows = []
    for code in ("LX", "CC", "AH"):
        eff = calc.EFFECTIVE[code]
        for fid, v in sorted(f["details"][code].items()):
            rows.append([code, NAMES[code], fid, pn.f[fid]["cls"], pn.f[fid]["region"], W.ym(eff), "2026-06",
                         "%.2f" % v["actual"], "%.2f" % v["reference"], "%.2f" % v["tax_effect"],
                         ";".join(v["peers"])])
    with open(G / CSVF, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["change_code", "change", "filer_id", "class", "region", "window_from", "window_to",
                    "actual_taxable_receipts", "reference_taxable_receipts", "tax_effect", "reference_filers"])
        w.writerows(rows)


# ---------------------------------------------------------------------------------------------- the bridge chart
def bridge_png(f):
    order = ["LX", "CC", "CO", "AH", "PC"]
    fc, actual = f["forecast_fy2026"], f["actual_fy2026"]
    short = {"LX": "Long-stay\nexemption", "CC": "Convention\ncenter closure", "CO": "County\noption cut",
             "AH": "Airline hub\nreduction", "PC": "Facilitator\ncollection"}
    steps = [("Adopted\nforecast", fc, "total")] + [(short[c], f["effects"][c], "step") for c in order]
    steps += [("No change\naccounts for", f["residual"], "step"), ("Actual", actual, "total")]
    fig, ax = plt.subplots(figsize=(10.5, 4.6), dpi=200)
    run = 0.0
    xs = range(len(steps))
    lo_vals = []
    for i, (lab, v, kind) in enumerate(steps):
        if kind == "total":
            ax.bar(i, v / 1e6, color=NAVY, width=0.62)
            ax.text(i, v / 1e6 + 1.2, "{:,}".format(k(v)), ha="center", fontsize=8.5, color=NAVY, fontweight="bold")
            run = v
            lo_vals.append(v)
        else:
            new = run + v
            col = RED if v < 0 else GREEN
            if lab.startswith("Long-stay"):
                col = "#8E2C1A"
            ax.bar(i, abs(v) / 1e6, bottom=min(run, new) / 1e6, color=col, width=0.62)
            ax.text(i, max(run, new) / 1e6 + 1.2, ("+" if v > 0 else "-") + "{:,}".format(abs(k(v))), ha="center",
                    fontsize=8.5, color=col)
            ax.plot([i - 0.31 - 0.38, i - 0.31], [run / 1e6, run / 1e6], color=GREY, lw=0.6)
            run = new
            lo_vals.append(new)
    ax.set_xticks(list(xs))
    ax.set_xticklabels([s[0] for s in steps], fontsize=8)
    ax.set_ylim(min(lo_vals) / 1e6 - 12, fc / 1e6 + 6)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: "\\$%dM" % v))
    ax.set_ylabel("Lodging tax due, fiscal 2026 periods", fontsize=8)
    ax.tick_params(axis="y", labelsize=8)
    ax.spines[["top", "right"]].set_visible(False)
    ax.set_title("Of the five changes, the long-stay exemption cost the most: \\$%s thousand of a \\$%s thousand shortfall"
                 % ("{:,}".format(abs(k(f["effects"]["LX"]))), "{:,}".format(k(f["shortfall"]))),
                 fontsize=10, loc="left", fontweight="bold")
    ax.text(0.0, -0.2, "Thousands of dollars. Each change measured under the standing procedure; each figure rounded "
                       "once, so the bars need not foot to the dollar.", transform=ax.transAxes, fontsize=7,
            color="#555555")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------------------------- the deck
def _text(slide, x, y, w, h, paras, size=14, bold_first=False, color="222222"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, t in enumerate(paras):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = par.add_run()
        run.text = t
        run.font.size = Pt(size)
        run.font.name = "Calibri"
        run.font.bold = bold_first and i == 0
        run.font.color.rgb = RGBColor.from_string(color)
        par.space_after = Pt(6)
    return tb


def _title(slide, text, sub=None):
    _text(slide, 0.5, 0.3, 12.3, 0.9, [text], size=24, bold_first=True, color="2F4858")
    if sub:
        _text(slide, 0.5, 1.05, 12.3, 0.4, [sub], size=12, color="555555")


def _footer(slide, n):
    _text(slide, 0.5, 7.05, 9, 0.3, ["Office of Revenue Analysis  |  Fiscal 2026 lodging tax variance statement"],
          size=9, color="777777")
    _text(slide, 12.0, 7.05, 1, 0.3, [str(n)], size=9, color="777777")


def _table(slide, x, y, w, rows, widths, size=12, bold_last=False):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                   Inches(0.36 * len(rows)))
    tbl = shape.table
    for j, wd in enumerate(widths):
        tbl.columns[j].width = Inches(wd)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = val
            par = cell.text_frame.paragraphs[0]
            par.alignment = PP_ALIGN.RIGHT if j > 0 else PP_ALIGN.LEFT
            for run in par.runs:
                run.font.size = Pt(size)
                run.font.name = "Calibri"
                run.font.bold = i == 0 or (bold_last and i == len(rows) - 1)
                run.font.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else "222222")
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string("2F4858" if i == 0 else ("F2F4F6" if i % 2 else "FFFFFF"))
    return tbl


def deck(f, d):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    eff = f["effects"]
    lx, cc = abs(k(eff["LX"])), abs(k(eff["CC"]))

    s = prs.slides.add_slide(blank)
    _title(s, "Fiscal 2026 lodging tax: the long-stay exemption is the change that cost the most",
           "Variance statement for the Joint Committee on Finance, hearing of September 15, 2026")
    _text(s, 0.5, 1.9, 12.3, 4.5, [
        "The statement names the long-stay exemption. Measured under the standing procedure it reduced fiscal 2026 "
        "lodging tax by $%s thousand, more than any other change of the year." % "{:,}".format(lx),
        "Tax due on fiscal 2026 returns came in $%s thousand under the adopted forecast. The five changes account for "
        "$%s thousand of that; the remaining $%s thousand is not accounted for by any change of the year." % (
            "{:,}".format(k(f["shortfall"])), "{:,}".format(abs(k(sum(eff.values())))),
            "{:,}".format(abs(k(f["residual"])))),
        "The convention center closure looks larger on a plain comparison because downtown hotels fell hardest. Their "
        "matched reference hotels, the ones with the same pattern through the year, fell nearly as far, so the "
        "closure itself accounts for $%s thousand." % "{:,}".format(cc),
        "Before it was used for this statement the procedure reproduced all %d reference figures in Appendix B of the "
        "fiscal 2024 and 2025 statements, to the dollar." % f["reproduced"][0]], size=16)
    _footer(s, 1)

    s = prs.slides.add_slide(blank)
    _title(s, "Forecast to actual: the long-stay exemption is the largest change; most of the gap is no change at all")
    s.shapes.add_picture(bridge_png(f), Inches(0.5), Inches(1.2), width=Inches(12.3))
    _footer(s, 2)

    s = prs.slides.add_slide(blank)
    _title(s, "Each change's effect under the standing procedure", "Thousands of dollars, fiscal 2026 periods")
    rows = [["", "Effect on fiscal 2026 tax"]]
    for c in sorted(eff, key=lambda c: eff[c]):
        rows.append([NAMES[c], ("+" if eff[c] > 0 else "-") + "{:,}".format(abs(k(eff[c])))])
    rows.append(["Not accounted for by any change", "-" + "{:,}".format(abs(k(f["residual"])))])
    rows.append(["Adopted forecast", "{:,}".format(k(f["forecast_fy2026"]))])
    rows.append(["Tax due on fiscal 2026 returns", "{:,}".format(k(f["actual_fy2026"]))])
    rows.append(["Shortfall against forecast", "-" + "{:,}".format(k(f["shortfall"]))])
    _table(s, 0.5, 1.7, 8.0, rows, [5.6, 2.4], size=13, bold_last=True)
    _text(s, 8.9, 1.7, 3.9, 4.6, [
        "The closure and the airline cut are measured filer by filer against reference hotels with the same monthly "
        "pattern in the year before the change, outside their own region.",
        "The county option cut is measured at the old and new rates; facilitator collection is measured on hosts and "
        "facilitators together.",
        "Per-filer reference figures are in fy2026_lodging_reference_matches.csv."], size=12, color="444444")
    _footer(s, 3)

    cr = f["cash"]
    s = prs.slides.add_slide(blank)
    _title(s, "Reconciling the treasury's cash collections to the statement",
           "Thousands of dollars. The monthly report counts deposits booked in the year; the statement counts tax "
           "due on the year's periods")
    rows = [["", "Fiscal 2026"],
            ["Tax due on fiscal 2026 returns (statement)", "{:,}".format(k(cr["liability"]))],
            ["Less fiscal 2026 periods not deposited by June 30", "-" + "{:,}".format(k(cr["own_periods_not_deposited"]))],
            ["Plus earlier periods deposited in fiscal 2026", "+" + "{:,}".format(k(cr["earlier_periods_deposited"]))],
            ["Plus amended returns booked in fiscal 2026", ("+" if cr["amendments"] >= 0 else "-")
             + "{:,}".format(abs(k(cr["amendments"])))],
            ["Less refunds booked in fiscal 2026", "-" + "{:,}".format(abs(k(cr["refunds"])))],
            ["Treasury collections, fiscal 2026", "{:,}".format(k(cr["cash"]))]]
    _table(s, 0.5, 1.9, 9.0, rows, [6.6, 2.4], size=13, bold_last=True)
    _text(s, 0.5, 5.0, 12.3, 1.2, ["Collections exclude penalty and interest. A second amendment to a return replaces "
                                   "the first. Figures rounded once each, so the column need not foot to the dollar."],
          size=11, color="555555")
    _footer(s, 4)

    s = prs.slides.add_slide(blank)
    _title(s, "Sources and method")
    _text(s, 0.5, 1.5, 12.3, 5, [
        "Returns extract, fiscal 2023 to 2026, and filer register, Department of Revenue, as of August 14, 2026.",
        "Treasury lodging tax ledger, July 2024 to July 2026. Adopted forecast: December 2024 consensus.",
        "Standing procedure, revision of March 2023. Reference filers: no change of the year, outside the filer's "
        "own region, a return in each of the 24 months to June 2026, the five closest on the year-before monthly "
        "pattern (each month's receipts standardised against the filer's own average and spread), no reference filer "
        "used for more than two touched filers, larger touched filers matched first, each month carried forward at "
        "the median growth of its reference filers.",
        "Prepared by %s, Senior Economist, for %s, Director." % (W.PEOPLE["analyst"], W.PEOPLE["director"])],
          size=13, color="333333")
    _footer(s, 5)

    prs.core_properties.author = W.PEOPLE["analyst"]
    prs.core_properties.last_modified_by = W.PEOPLE["director"]
    prs.core_properties.title = "Fiscal 2026 lodging tax variance statement"
    prs.core_properties.revision = 4
    prs.core_properties.created = dt.datetime(2026, 9, 2, 14, 10)
    prs.core_properties.modified = dt.datetime(2026, 9, 8, 21, 45)
    prs.core_properties.comments = ""
    prs.save(G / DECK)


def scrub(path, when):
    src = zipfile.ZipFile(path)
    items = [(i, src.read(i.filename)) for i in src.infolist()]
    src.close()
    app = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
           '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
           '<Application>Microsoft Office PowerPoint</Application></Properties>').encode()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            if info.filename == "docProps/thumbnail.jpeg":
                continue
            if info.filename == "docProps/app.xml":
                data = app
            elif info.filename == "_rels/.rels":
                data = re.sub(rb'<Relationship [^>]*Target="docProps/thumbnail.jpeg"[^>]*/>', b"", data)
            elif info.filename == "[Content_Types].xml":
                data = re.sub(rb'<Default Extension="jpeg"[^>]*/>', b"", data)
            zi = zipfile.ZipInfo(info.filename, date_time=when.timetuple()[:6])
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, data)
    path.write_bytes(buf.getvalue())


def build():
    G.mkdir(exist_ok=True)
    for p in G.iterdir():
        p.unlink()
    f, d = figures()
    write_csv(f, d)
    deck(f, d)
    import seal
    scrub(G / DECK, seal.DECK_SENT)
    seal.stamp_dir(G, seal.DECK_SENT)
    return f


if __name__ == "__main__":
    f = build()
    print({c: k(v) for c, v in f["effects"].items()}, f["named"], k(f["shortfall"]), k(f["residual"]), f["reproduced"])
