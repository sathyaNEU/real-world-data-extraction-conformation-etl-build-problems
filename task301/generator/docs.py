"""The documents: the executive paper that sets the 2027 placement, and the PHN's data dictionary for the folder."""
import datetime as dt

from docx import Document
from docx.shared import Pt

import params as P
import postprocess

AUTHOR = "Megan Collins"


def _doc():
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    cp = d.core_properties
    cp.author = AUTHOR
    cp.last_modified_by = AUTHOR
    cp.created = dt.datetime(2026, 10, 14, 9, 12)
    cp.modified = dt.datetime(2026, 10, 20, 16, 41)
    cp.revision = 4
    cp.title = ""
    cp.comments = ""
    return d


def plan(path):
    d = _doc()
    d.add_heading("Quit outreach team: placement for 2027", 1)
    d.add_paragraph("Executive meeting, 11 November 2026. Item 6. Prepared by Health Intelligence.")
    d.add_heading("Background", 2)
    d.add_paragraph(
        "The quit outreach team is a clinical nurse consultant and two quit coaches. The Commonwealth's 2026-27 "
        "allocation funds it again from 1 January 2027 for one calendar year, in one of four sub-regions: Coastal "
        "Hinterland, Inner South, Northern Corridor or Western Plains.")
    d.add_paragraph(
        "The team worked in Eastern Shore in 2025. It telephoned the active patients Eastern Shore's practices "
        "recorded as current smokers at the end of January 2025, making up to three attempts to reach each one between "
        "February and November, and offered each smoker it reached a quit date with nurse and coach support. Its case "
        "log records the outcome for every patient on its list and, for each patient who set a quit date, whether they "
        "were smoking four weeks after that date.")
    d.add_heading("How the team will work in 2027", 2)
    d.add_paragraph(
        "The team will call every active patient the chosen sub-region's practices record as a current smoker in the "
        "PHN's 30 September 2026 extract, once each, spreading the calls evenly through 2027. It works as it did in "
        "Eastern Shore: the same staff, the same three attempts, the same script and the same support after a quit date.")
    d.add_heading("The decision", 2)
    d.add_paragraph(
        "The executive will place the team in the sub-region where it will help the most patients to quit in 2027. A "
        "quit is a patient the team calls in 2027 who sets a quit date with the team and is not smoking four weeks "
        "after that date. A patient the team cannot contact at four weeks is counted as smoking, as the state's quit "
        "services count them.")
    d.add_heading("The data", 2)
    d.add_paragraph(
        "The active patient extract holds every active patient aged 18 or over at the 30 September 2026 extract, with "
        "the most recent smoking status their practice has recorded and the date it was recorded. The status history "
        "holds every smoking status those practices have recorded for those patients since 1 January 2012, as the "
        "extraction tool took it from each practice's clinical system. The data dictionary describes both, the team's "
        "case log and the other files in the folder.")
    d.add_paragraph(
        "Health Intelligence will bring the placement, with the Quitline referrals from each candidate sub-region's "
        "practices and what the 2025 team's calls show about evening work and talk time, to the meeting.")
    d.save(path)
    postprocess.scrub(path, P.ORG, "2026-10-20")
    postprocess.finish_office(path, "2026-10-14T09:12:00Z", "2026-10-20T16:41:00Z", P.ORG, person=AUTHOR, revision=4)


def dictionary(path):
    d = _doc()
    d.core_properties.created = dt.datetime(2026, 3, 2, 10, 5)
    d.core_properties.modified = dt.datetime(2026, 10, 15, 11, 30)
    d.add_heading("Primary care data folder: data dictionary", 1)
    d.add_paragraph("Coast and Tablelands PHN, Health Intelligence. Version 2.3, October 2026.")

    def table(rows):
        t = d.add_table(rows=1, cols=2)
        t.style = "Light Grid Accent 1"
        t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Field", "Description"
        for a, b in rows:
            r = t.add_row().cells
            r[0].text, r[1].text = a, b

    d.add_heading("active_patients_2026-09-30.csv", 2)
    d.add_paragraph("One row per active patient aged 18 or over at the 30 September 2026 extract. A patient is active "
                    "when they have attended the practice three or more times in the past two years.")
    table([("patient_key", "The patient's key, the same in every file in the folder."),
           ("practice_id", "The practice, as in practices.csv."),
           ("sub_region", "The PHN sub-region the practice belongs to."),
           ("age_band", "Age at the extract, in five-year bands."),
           ("sex", "As recorded by the practice."),
           ("first_seen", "The date the practice first saw the patient."),
           ("chronic_condition_review", "Y when the patient has had a chronic disease management plan or review in "
                                        "the past two years."),
           ("smoking_status", "The most recent smoking status recorded, in the wording of the practice's clinical "
                              "system. Blank when none has been recorded."),
           ("smoking_status_date", "The date that status was recorded.")])
    d.add_heading("smoking_status_history_2012-2026.parquet", 2)
    d.add_paragraph("One row per smoking status entry since 1 January 2012 for the patients in the extract, in "
                    "the wording of the clinical system that recorded it.")
    table([("patient_key", "As in the extract."),
           ("practice_id", "The practice holding the record at the extract."),
           ("recorded_at", "The date and time the entry was saved."),
           ("smoking_status", "The status entered.")])
    d.add_paragraph(
        "Best Practice records a current smoker as a daily, weekly or irregular smoker and a person who has never "
        "smoked as never smoked; Medical Director records a current smoker as a smoker and a person who has never "
        "smoked as a non-smoker. Both record an ex-smoker as ex-smoker. A current smoker is any of daily, weekly or "
        "irregular smoker or smoker. A practice that changed clinical system brought its records across with their "
        "original dates and wording.")
    d.add_heading("practices.csv", 2)
    table([("practice_id", "The PHN's practice identifier."),
           ("practice_name", "The practice's trading name in the PHN register."),
           ("sub_region", "The PHN sub-region the practice belongs to."),
           ("clinical_system", "The clinical system at the extract."),
           ("previous_system, system_changed", "The earlier system and the date of the change, where a practice "
                                               "changed system after 2012."),
           ("active_patients_18_plus", "Active patients aged 18 or over at the extract.")])
    d.add_heading("quit_outreach_calls_eastern_shore_2025.csv", 2)
    d.add_paragraph("The 2025 team's case log: one row per patient on its list.")
    table([("patient_key, practice_id", "As in the extract. Patients who have since left Eastern Shore's practices "
                                        "are not in the 2026 extract."),
           ("list_status, list_status_date", "The status and date that put the patient on the list."),
           ("call_date", "The date of the team's last attempt, or of the call that reached the patient."),
           ("attempts", "Attempts made, up to three."),
           ("outcome", "Not reached; Not a current smoker (the patient told the team they no longer smoke); Declined "
                       "support; Quit date set."),
           ("quit_date", "The quit date agreed with the patient."),
           ("four_week_status", "Not smoking, Smoking or Lost to follow-up, from the team's call four weeks after the "
                                "quit date.")])
    d.add_heading("quitline_referrals_2023-2025.csv", 2)
    d.add_paragraph("The state Quitline's monthly report to the PHN of referrals it received from general practices "
                    "in the PHN's area.")
    table([("month", "The month Quitline received the referrals."),
           ("referring_practice", "The referring practice, as Quitline's directory names it."),
           ("phn_sub_region", "The PHN sub-region in Quitline's report."),
           ("channel", "eReferral or Fax."),
           ("referrals", "Referrals received.")])
    d.add_heading("dialler_legs_2025.csv", 2)
    d.add_paragraph("The team's dialler export for 2025: one row per leg the dialler logged.")
    table([("leg_id", "The dialler's row identifier."),
           ("patient_key", "As in the case log."),
           ("agent_id", "The team member's dialler login, as in team_roster_2025.csv."),
           ("started", "The dialler's time stamp for the start of the leg."),
           ("duration_s", "Seconds from the start of the leg until the line was released."),
           ("result", "Answered, No answer, Voicemail or Busy.")])
    d.add_heading("team_roster_2025.csv", 2)
    d.add_paragraph("The team's 2025 shifts, in local time.")
    d.add_heading("Other files", 2)
    d.add_paragraph("sub_region_profiles_2025.csv holds each sub-region's residents aged 18 or over, remoteness and "
                    "median IRSD decile. gp_workforce_2025.xlsx holds each practice's GP and practice nurse "
                    "full-time equivalents and whether it is taking new patients. "
                    "quit_medicines_dispensed_2025.csv holds the PBS varenicline and nicotine replacement dispensings "
                    "for residents of each sub-region by month.")
    d.save(path)
    postprocess.scrub(path, P.ORG, "2026-10-15")
    postprocess.finish_office(path, "2026-03-02T10:05:00Z", "2026-10-15T11:30:00Z", P.ORG, person=AUTHOR, revision=9)
