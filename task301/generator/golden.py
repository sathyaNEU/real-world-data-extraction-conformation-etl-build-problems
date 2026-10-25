"""The two goldens: the placement paper (docx, with its chart) and the placement workbook (xlsx)."""
import datetime as dt
import io
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import params as P
import postprocess
import solve as S

DOCX = "quit_outreach_placement_2027.docx"
XLSX = "quit_outreach_placement_2027.xlsx"
AUTHOR = "Megan Collins"
INK, MUTED, ACCENT, BASE = "#1f2933", "#52606d", "#1d6fa5", "#b8c4ce"


def c(x):
    return "{:,}".format(int(x))


def chart(F, X):
    gold = X["gold"]
    fig, ax = plt.subplots(figsize=(6.4, 3.2), dpi=150)
    names = S.CANDS
    vals = [gold[g]["quits"] for g in names]
    cols = [ACCENT if g == F["call_region"] else BASE for g in names]
    ax.barh(range(len(names)), vals, color=cols, height=0.56)
    ax.set_yticks(range(len(names)))
    ax.set_yticklabels(names, color=INK, fontsize=9)
    ax.invert_yaxis()
    i = names.index(F["call_region"])
    ax.text(vals[i] + 8, i, "about %d" % F["call_quits"], va="center", fontsize=9, color=INK)
    ax.set_xlabel("Expected quits in 2027", color=MUTED, fontsize=9)
    ax.set_xlim(0, max(vals) * 1.18)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.spines["left"].set_color("#9aa5b1")
    ax.spines["bottom"].set_color("#9aa5b1")
    ax.tick_params(axis="x", colors=MUTED, labelsize=8)
    ax.grid(axis="x", color="#e4e7eb", linewidth=0.6)
    ax.set_axisbelow(True)
    ax.set_title("Quits the team would help in 2027, by sub-region", fontsize=10, color=INK, loc="left")
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", metadata={"Software": None})
    plt.close(fig)
    buf.seek(0)
    return buf


def paper(path, F, X):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    cp = d.core_properties
    cp.author = AUTHOR
    cp.last_modified_by = AUTHOR
    cp.created = dt.datetime(2026, 10, 27, 8, 50)
    cp.modified = dt.datetime(2026, 11, 3, 15, 22)
    cp.revision = 6
    cp.title = ""
    cp.comments = ""
    p = F["pilot"]
    win = F["call_region"]
    d.add_heading("Quit outreach team: 2027 placement", 1)
    d.add_paragraph("Executive meeting, 11 November 2026. Item 6. Health Intelligence.")
    d.add_heading("Recommendation", 2)
    d.add_paragraph("Place the quit outreach team in %s for 2027. On the 2025 team's results it will help about %d "
                    "patients there to quit, more than in any of the other three sub-regions." % (win, F["call_quits"]))
    d.add_heading("Why %s" % win, 2)
    d.add_paragraph(
        "The team can only help a patient who still smokes when it calls. A recorded status is a snapshot: some "
        "recorded smokers have stopped since, and how many depends on how long ago the status was recorded and on how "
        "fast smokers in that sub-region stop. Each sub-region's own records answer the second question, because the "
        "practices re-record many patients' status over the years, and they show smokers in Coastal Hinterland "
        "stopping more than twice as fast as smokers in %s." % win)
    d.add_paragraph(
        "Coastal Hinterland has the longest list, %s recorded smokers, but %s of them were recorded more than five years "
        "before the extract, and at its rate most of those no longer smoke. %s's %s recorded smokers are mostly recent, "
        "and about %d per cent of them will still smoke when the team calls in 2027."
        % (c(F["listed"]["Coastal Hinterland"]), c(F["listed_over5y"]["Coastal Hinterland"]), win, c(F["listed"][win]),
           F["call_still_pct"]))
    d.add_paragraph(
        "In Eastern Shore in 2025 the team reached %s of the %s patients on its list (%.1f per cent); %s of those told it "
        "they had already stopped, and %s of the smokers it reached (%.1f per cent) set a quit date and were not smoking "
        "four weeks later. That is about %d per cent of the listed patients who still smoked when called, which the "
        "estimate carries to 2027 because the team will work the same way."
        % (c(p["reached"]), c(p["listed"]), F["pilot_reach_pct"], c(p["stopped"]), c(p["quits"]),
           F["pilot_quit_per_smoker_reached_pct"], F["quits_per100_still"]))
    d.add_picture(chart(F, X), width=Cm(15.5))
    d.add_heading("Points for the executive", 2)
    d.add_paragraph("Eastern Shore's 2025 list was unusually fresh after its practices' 2024 recording drive, so its "
                    "quits per listed patient overstate what a list with older records gives.", style="List Bullet")
    d.add_paragraph("A patient the team cannot contact at four weeks is counted as smoking, as the plan sets out.",
                    style="List Bullet")
    d.add_paragraph("Quitline referrals, evening calling and talk time are in the workbook.", style="List Bullet")
    d.save(path)
    postprocess.scrub(path, P.ORG, "2026-11-03")
    postprocess.finish_office(path, "2026-10-27T08:50:00Z", "2026-11-03T15:22:00Z", P.ORG, person=AUTHOR, revision=6)


def workbook(path, F, X):
    wb = Workbook()
    hdr = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1D6FA5")
    ws = wb.active
    ws.title = "Placement"
    ws.append(["Quit outreach team: 2027 placement"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Recommended sub-region", F["call_region"]])
    ws.append(["Expected quits in 2027 (nearest 50)", F["call_quits"]])
    ws.append([])
    ws.append(["Sub-region", "Recorded current smokers, 30 Sep 2026", "Recorded more than five years earlier"])
    for cell in ws[5]:
        cell.font, cell.fill = hdr, fill
    for g in S.CANDS:
        ws.append([g, F["listed"][g], F["listed_over5y"][g]])
    ws.append([])
    ws.append(["2025 team in Eastern Shore", "Patients"])
    for k, lab in (("listed", "On the list"), ("reached", "Reached"), ("stopped", "Already stopped"),
                   ("quit_dates", "Set a quit date"), ("quits", "Not smoking at four weeks")):
        ws.append([lab, F["pilot"][k]])
    ws.column_dimensions["A"].width = 34
    for col in "BC":
        ws.column_dimensions[col].width = 34
    ws2 = wb.create_sheet("Quitline referrals")
    ws2.append(["Quitline referrals from each candidate sub-region's practices, 2025"])
    ws2["A1"].font = Font(bold=True)
    ws2.append(["Sub-region", "Referrals 2025"])
    for cell in ws2[2]:
        cell.font, cell.fill = hdr, fill
    for g_ in S.CANDS:
        ws2.append([g_, F["ask_referrals"][g_]])
    ws2.append([])
    ws2.append(["Practices counted in the sub-region the PHN register gives them."])
    ws2.column_dimensions["A"].width = 24
    ws3 = wb.create_sheet("Evening calls")
    ws3.append(["2025 team: calls to patients from 5 pm, by weekday"])
    ws3["A1"].font = Font(bold=True)
    ws3.append(["Weekday", "Calls from 5 pm", "Patients reached"])
    for cell in ws3[2]:
        cell.font, cell.fill = hdr, fill
    for d in S.WEEKDAYS:
        ws3.append([d, F["ask_evening"][d]["calls"], F["ask_evening"][d]["reached"]])
    ws3.column_dimensions["A"].width = 14
    ws3.column_dimensions["B"].width = 16
    ws3.column_dimensions["C"].width = 18
    ws4 = wb.create_sheet("Talk time")
    ws4.append(["2025 team: minutes on answered calls per patient reached, by team member's dialler login"])
    ws4["A1"].font = Font(bold=True)
    ags = list(F["ask_talk_min"])
    ws4.append(["Month"] + ags)
    for cell in ws4[2]:
        cell.font, cell.fill = hdr, fill
    for m in F["ask_talk_min"][ags[0]]:
        ws4.append([m] + [F["ask_talk_min"][ag][m] for ag in ags])
        for j in range(2, 2 + len(ags)):
            ws4.cell(row=ws4.max_row, column=j).number_format = "0.0"
    ws4.column_dimensions["A"].width = 14
    for col in ("B", "C", "D"):
        ws4.column_dimensions[col].width = 12
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.properties.created = dt.datetime(2026, 10, 27, 9, 5)
    wb.properties.modified = dt.datetime(2026, 11, 3, 14, 58)
    wb.save(path)
    postprocess.scrub(path, P.ORG, "2026-11-03")
    postprocess.finish_office(path, "2026-10-27T09:05:00Z", "2026-11-03T14:58:00Z", P.ORG, person=AUTHOR)


def write(folder, F, X):
    os.makedirs(folder, exist_ok=True)
    paper(os.path.join(folder, DOCX), F, X)
    workbook(os.path.join(folder, XLSX), F, X)
