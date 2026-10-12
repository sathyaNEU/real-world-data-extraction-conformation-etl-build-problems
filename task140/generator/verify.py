"""Independent verifier: reads only the shipped pack (target/) and the goldens, on its own code path (pandas), and
recomputes the guarantee, the default path's guarantee and every 2026 Q3 table, then compares with the golden
workbook. Shares no code with calc.py."""
from __future__ import annotations

import datetime as dt
import json
import math
import os
import re
import sys

import numpy as np
import pandas as pd


def half_up(x, step):
    return int(math.floor(x / step + 0.5)) * step


def find(target, pattern):
    hits = [f for f in os.listdir(target) if re.search(pattern, f)]
    assert len(hits) == 1, (pattern, hits)
    return os.path.join(target, hits[0])


def main_call(target):
    dirx = json.load(open(find(target, r"^tenant_directory\.json$"), encoding="utf-8"))
    ten = pd.DataFrame(dirx["tenants"])
    firm = ten[ten["class"] == "firm"].copy()
    firm["code"] = firm["cost_centre"].str.extract(r"CC-(\d+)")[0]
    plan = pd.read_excel(find(target, r"hardware_plan"), sheet_name="By quarter", header=None)
    plan = plan[plan[0].astype(str).str.fullmatch(r"\d{4}Q\d")]
    machines = int(plan[plan[0] == "2027Q1"][3].iloc[0])
    import docx
    memo = " ".join(p.text for p in docx.Document(find(target, r"capacity_memo")).paragraphs)
    oc = float(re.search(r"overcommit for 2027 Q1 is ([\d.]+)", memo).group(1))
    reservable = round(machines * 96 * oc)
    req = pd.read_csv(find(target, r"capacity_requests"))
    bud = pd.read_excel(find(target, r"tech_budgets"), sheet_name="Compute budgets", header=None)
    hdr = bud.index[bud[0] == "CC"][0]
    bud.columns = bud.loc[hdr]
    bud = bud.loc[hdr + 1:]
    bud = bud[bud["Budget line"] == "MAD1 C-22 reserved cores"].copy()
    bud["code"] = bud["CC"].astype(str)
    rc = pd.read_csv(find(target, r"rate_card"))
    rc = rc[(rc.resource == "reserved core") & (rc.site == "MAD1")].sort_values("effective_from")
    rate = float(rc[rc.effective_from <= "2027-01-01"].iloc[-1].eur_per_unit_month)
    m = req.merge(firm[["id", "code"]], left_on="tenant", right_on="id").merge(bud[["code", "Jan 2027"]], on="code")
    m["cap"] = np.floor(m["Jan 2027"].astype(float) / rate + 1e-9).astype(int)
    m["approved"] = np.minimum(m["cores_requested"], m["cap"])
    F = int(m["approved"].sum())
    use = pd.read_csv(find(target, r"usage_hourly")).drop_duplicates()
    wide = use.pivot(index="hour_start_utc", columns="tenant", values="cores_in_use")
    cell = pd.read_csv(find(target, r"cell_hourly")).set_index("hour_start_utc")
    df = wide.join(cell, how="inner").dropna(subset=["cell-store", "batch-train", "machines_in_service"])
    X = np.column_stack([np.ones(len(df)), df["machines_in_service"].values, df["batch-train"].values])
    a, c, b = np.linalg.lstsq(X, df["cell-store"].values, rcond=None)[0]
    q = pd.PeriodIndex(pd.to_datetime(df.index.str[:19]), freq="Q").astype(str)
    peaks = df.groupby(q)[["cell-ctl", "cell-store"]].max()
    last4 = peaks.iloc[-4:]
    x = np.arange(4)
    ctl = np.polyval(np.polyfit(x, last4["cell-ctl"].values, 1), 4)
    store_trend = np.polyval(np.polyfit(x, last4["cell-store"].values, 1), 4)
    Gu = (reservable - F - ctl - a - c * machines) / (1 + b)
    G = int(Gu // 1000) * 1000
    S = a + c * machines + b * G
    default = int((reservable - F - ctl - store_trend) // 1000) * 1000
    return dict(G=G, F=F, S=S, ctl=ctl, default=default, Gu=Gu)


def asks(target):
    runs = pd.read_csv(find(target, r"batch_runs"), dtype=str, keep_default_na=False)
    ck = pd.read_csv(find(target, r"batch_checkpoints"), dtype=str, keep_default_na=False).set_index("checkpoint_id")
    rz = pd.read_csv(find(target, r"batch_resizes"), dtype={"run_id": str})
    q0, q1 = pd.Timestamp("2026-07-01", tz="UTC"), pd.Timestamp("2026-10-01", tz="UTC")
    ts = lambda s: pd.Timestamp(s) if s else q1
    rz["at"] = pd.to_datetime(rz["at_utc"])
    rzg = {k: list(zip(v["at"], v["cores_after"])) for k, v in rz.sort_values("at").groupby("run_id")}

    def ch(row, a, b):
        pts = [(ts(row.started_utc), int(row.cores))] + rzg.get(row.run_id, [])
        end = ts(row.ended_utc)
        tot = 0.0
        for i, (t, v) in enumerate(pts):
            nt = pts[i + 1][0] if i + 1 < len(pts) else end
            lo, hi = max(t, a), min(nt, b, end)
            if hi > lo:
                tot += v * (hi - lo).total_seconds() / 3600
        return tot
    nxt = runs[runs.resumed_from != ""].set_index("resumed_from")
    held, red = {}, {}
    for row in runs.itertuples():
        held[row.team] = held.get(row.team, 0.0) + ch(row, q0, q1)
        if row.end_state == "PREEMPTED" and q0 <= ts(row.ended_utc) < q1:
            c = ck.loc[nxt.loc[row.run_id, "restore_checkpoint"]]
            t0 = ts(c.write_started_utc) if c.run_id == row.run_id else ts(row.started_utc)
            red[row.team] = red.get(row.team, 0.0) + ch(row, max(t0, q0), q1)
    A = {t: (half_up(red.get(t, 0.0), 100), math.floor(10000 * red.get(t, 0.0) / held[t] + 0.5) / 100) for t in held}
    ev = pd.read_csv(find(target, r"machine_events")).drop_duplicates()
    ev["t"] = pd.to_datetime(ev["time_utc"])
    ids = pd.read_excel(find(target, r"hardware_inventory"), sheet_name="Cluster ids", header=None)
    ids = ids[ids[0].astype(str).str.startswith("SRV-")]
    asset = dict(zip(ids[1].astype(int), ids[0]))
    ev["asset"] = ev["machine_id"].map(asset)
    ev = ev.sort_values(["t", "event"])
    B = {}
    for _, grp in ev.groupby("asset"):
        rows = grp.sort_values("t").to_dict("records")
        for i, r in enumerate(rows):
            if r["event"] != "REMOVED":
                continue
            back = next((x for x in rows[i + 1:] if x["event"] == "REGISTERED"), None)
            if back is not None:
                k = (r["cause"], r["t"].month)
                B[k] = B.get(k, 0.0) + (back["t"] - r["t"]).total_seconds() / 3600
    B = {k: half_up(v, 10) for k, v in B.items()}
    led = pd.read_csv(find(target, r"reservation_ledger"))
    base = led[(led.quarter == "2026Q3") & (led["class"] == "firm")].set_index("tenant")["reserved_cores"]
    chg = pd.read_csv(find(target, r"quota_changes"))
    qr = pd.read_csv(find(target, r"reconciler_runs"), keep_default_na=False)
    app = {}
    for r in qr.itertuples():
        for cid in [x for x in r.changes_applied.split(";") if x]:
            app[cid] = pd.Timestamp(r.run_started_utc)
    chg["applied"] = chg["change_id"].map(app)
    rc = pd.read_csv(find(target, r"rate_card"))
    rc = rc[(rc.resource == "reserved core") & (rc.site == "MAD1")]
    Cc = {}
    for t, lvl in base.items():
        tot = 0.0
        for d in pd.date_range("2026-07-01", "2026-09-30", freq="D", tz="UTC"):
            h = lvl + chg[(chg.tenant == t) & (chg.applied <= d)]["delta_cores"].sum()
            rate = float(rc[rc.effective_from <= d.strftime("%Y-%m-%d")].sort_values("effective_from")
                         .iloc[-1].eur_per_unit_month)
            tot += h * rate / d.days_in_month
        Cc[t] = half_up(tot, 100)
    return A, B, Cc


def main(target, golden):
    import openpyxl
    mc = main_call(target)
    wb = openpyxl.load_workbook(os.path.join(golden, [f for f in os.listdir(golden) if f.endswith(".xlsx")][0]),
                                data_only=True)
    top = {r[0]: r for r in wb["2027 Q1"].iter_rows(values_only=True) if r and isinstance(r[0], str)}
    assert mc["G"] == top["Batch training guarantee"][2], (mc["G"], top["Batch training guarantee"])
    assert half_up(mc["F"], 1000) == top["Firm tenants (eleven)"][2]
    assert half_up(mc["S"], 1000) == top["cell-store"][2] and half_up(mc["ctl"], 1000) == top["cell-ctl"][2]
    assert mc["default"] != mc["G"]
    A, B, Cc = asks(target)
    sa = {r[0]: r for r in wb["Training lost 2026 Q3"].iter_rows(values_only=True) if r and r[0] in A}
    for t, (red, share) in A.items():
        assert sa[t][2] == red and abs(sa[t][3] - share) < 1e-9, (t, sa[t], red, share)
    sb = {r[0]: r for r in wb["Machine-hours lost 2026 Q3"].iter_rows(values_only=True)
          if r and r[0] in ("July", "August", "September")}
    for (cause, mth), v in B.items():
        lab = {7: "July", 8: "August", 9: "September"}[mth]
        assert sb[lab][1 if cause == "hw_fault" else 2] == v, (cause, mth, v, sb[lab])
    sc = {r[0]: r for r in wb["Chargeback 2026 Q3"].iter_rows(values_only=True) if r and r[0] in Cc}
    for t, v in Cc.items():
        assert sc[t][2] == v, (t, v, sc[t])
    print("verify: guarantee %s, default path %s, firm %s, cell-store %.1f, cell-ctl %.1f; Ask A %d teams, Ask B %d "
          "figures, Ask C %d tenants agree with the golden workbook" % (
              "{:,}".format(mc["G"]), "{:,}".format(mc["default"]), "{:,}".format(mc["F"]), mc["S"], mc["ctl"],
              len(A), len(B), len(Cc)))
    return True


if __name__ == "__main__":
    t = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                                                           "target")
    g = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(t), "golden")
    main(t, g)
