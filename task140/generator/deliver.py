"""The golden deliverables: c22_capacity_2027q1.xlsx (the 2027 Q1 guarantee and the reservations it is struck after,
the firm approvals, the storage forecast, and the three 2026 Q3 tables) and c22_capacity_review_2027q1.pptx (one
slide: how the quarter's reservable cores divide between the firm tenants, the platform services and batch), plus
generator/figures.json. Every figure is computed from the shipped pack (calc.py)."""
from __future__ import annotations

import json
import math
import os

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import calc as C
import common as K
import postprocess as PP

MONTHS = {7: "July", 8: "August", 9: "September"}


def pct2(x):
    return "%.2f" % (math.floor(x * 100 + 0.5) / 100.0)


def compute(target):
    M = C.load(target)
    lad = C.ladder(M)
    g = lad["R3"]
    A = C.load_asks(target)
    redone, share, held = C.ask_a(A)
    lost = C.ask_b(A)
    owe = C.ask_c(A)
    approvals = C.firm_approvals(M)
    return dict(M=M, lad=lad, g=g, A=A, redone=redone, share=share, held=held, lost=lost, owe=owe,
                approvals=approvals)


def f0(x):
    return "{:,}".format(int(x))


def figures(c):
    g, lad, M = c["g"], c["lad"], c["M"]
    out = {
        "call": {"guarantee_cores": f0(g["G"]),
                 "firm_tenants_cores": f0(K.round_to(g["F"], 1000)),
                 "cell_store_cores": f0(K.round_to(g["S"], 1000)),
                 "cell_ctl_cores": f0(K.round_to(g["ctl"], 1000)),
                 "platform_services_cores": f0(K.round_to(g["platform"], 1000)),
                 "reservable_cores": f0(g["R"])},
        "steps": {
            "machines_q1": f0(M.M1), "cores_per_machine": f0(K.CORES), "overcommit": "%.2f" % M.overcommit,
            "reservable_cores": f0(M.R1),
            "rate_q1_eur": "%.2f" % M.rate_q1,
            "firm_requests_total": f0(sum(M.requests.values())),
            "firm_approved_total": f0(g["F"]),
            "firm_approved": {t: f0(v) for t, v in c["approvals"].items()},
            "firm_held_to_budget": sorted(t for t in M.firm if c["approvals"][t] < M.requests[t]),
            "store_fit_intercept": "%.1f" % g["a"], "store_fit_per_machine": "%.4f" % g["c"],
            "store_fit_per_batch_core": "%.5f" % g["beta"], "store_fit_hours": f0(g["n"]),
            "one_plus_per_batch_core": "%.5f" % (1 + g["beta"]),
            "store_fit_resid_sd": "%.1f" % g["resid_sd"],
            "cell_ctl_forecast": "%.1f" % g["ctl"],
            "guarantee_unrounded": "%.1f" % g["Gu"],
            "cell_store_at_guarantee": "%.1f" % g["S"],
            "platform_total": "%.1f" % g["platform"],
            "unreserved": "%.1f" % g["unreserved"],
            "cs3_sections": ["2.1", "3.2", "3.3", "4.1", "4.2", "5.1", "5.2"],
        },
        "default_path": {"guarantee_cores": f0(lad["R0"]["G"]), "cell_store_cores": "%.1f" % lad["R0"]["S"],
                         "guarantee_unrounded": "%.1f" % lad["R0"]["Gu"]},
        "rungs": {k: {"guarantee_cores": f0(v["G"]), "guarantee_unrounded": "%.1f" % v["Gu"]} for k, v in lad.items()},
        "ask_a": {t: {"lost_core_hours": f0(K.round_to(c["redone"][t], 100)), "lost_share_pct": pct2(c["share"][t]),
                      "held_core_hours": f0(K.round_to(c["held"][t], 100))} for t in c["A"].teams},
        "ask_b": {MONTHS[m]: {"hardware_faults_machine_hours": f0(K.round_to(c["lost"][("hw_fault", m)], 10)),
                              "maintenance_machine_hours": f0(K.round_to(c["lost"][("maintenance", m)], 10))}
                  for m in (7, 8, 9)},
        "ask_c": {t: f0(K.round_to(c["owe"][t], 100)) for t in c["A"].firm},
    }
    return out


THIN = Side(style="thin", color="B4BCC8")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


def _title(ws, title, sub):
    ws["A1"] = title
    ws["A1"].font = Font(name="Calibri", bold=True, size=13, color="1F4E79")
    ws["A2"] = sub
    ws["A2"].font = Font(name="Calibri", italic=True, size=9.5, color="595959")


def _head(ws, row, values, widths=None):
    for k, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=k, value=v)
        cell.font = Font(name="Calibri", bold=True)
        cell.fill = PatternFill("solid", fgColor="DCE6F0")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.border = BOX
    if widths:
        for k, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(k)].width = w


def _row(ws, row, values, fmts=None, bold=False):
    for k, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=k, value=v)
        cell.border = BOX
        cell.font = Font(name="Calibri", bold=bold)
        if fmts and fmts[k - 1]:
            cell.number_format = fmts[k - 1]
        if isinstance(v, str) and len(v) > 40:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def write_book(path, c, figs):
    g, M, lad = c["g"], c["M"], c["lad"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2027 Q1"
    _title(ws, "C-22 capacity for 2027 Q1: batch guarantee and reservations",
           "Prepared for the review of 15 December 2026 under CS-3 (rev. June 2026). Cores are reservable cores. "
           "Monitoring export to 13 December 2026 23:00 UTC.")
    _head(ws, 4, ["Item", "Cores", "Reserved (nearest 1,000)", "How it is set"], [34, 13, 16, 92])
    r = 5
    items = [
        ("Machines in service from 28 December 2026", M.M1, None, "Hardware plan rev. 11 December 2026 (261 racks of 20)"),
        ("Reservable cores", M.R1, None, "%s machines x 96 cores x overcommit %.2f (capacity memo)" % (
            f0(M.M1), M.overcommit)),
        ("Firm tenants (eleven)", g["F"], K.round_to(g["F"], 1000), "Each Q1 request held to what its frozen monthly "
         "budget buys at EUR %.2f per core-month; sheet Firm tenants" % M.rate_q1),
        ("cell-ctl", round(g["ctl"], 1), K.round_to(g["ctl"], 1000), "Busiest hour forecast: line through the four "
         "quarters of 2026, one quarter on"),
        ("cell-store", round(g["S"], 1), K.round_to(g["S"], 1000), "Busiest hour forecast from its drivers: %.1f + "
         "%.4f per machine in service + %.5f per batch core running, at %s machines and batch at its guarantee "
         "(sheet cell-store)" % (g["a"], g["c"], g["beta"], f0(M.M1))),
        ("Platform services", round(g["platform"], 1), K.round_to(g["platform"], 1000), "cell-ctl and cell-store"),
        ("Batch training guarantee", g["G"], g["G"], "Largest whole thousand left; solved with cell-store, which "
         "rises by %.5f cores for each batch core: (%s - firm - cell-ctl - %.1f - %.4f x %s) / %.5f = %s" % (
             g["beta"], f0(M.R1), g["a"], g["c"], f0(M.M1), 1 + g["beta"], "{:,.1f}".format(g["Gu"]))),
        ("Not reserved", round(g["unreserved"], 1), None, "Left after the guarantee (CS-3 5.2)"),
    ]
    for lab, v, rv, how in items:
        _row(ws, r, [lab, v, rv, how], [None, "#,##0" if isinstance(v, int) else "#,##0.0", "#,##0", None],
             bold=(lab == "Batch training guarantee"))
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="For comparison: carrying cell-store forward on its own last four busiest hours "
            "(%s cores) would leave %s cores, but every batch core added brings %.5f cores of cell-store load that "
            "history at a guarantee of 55,000 to 58,000 cores has not seen." % (
                "{:,.0f}".format(lad["R0"]["S"]), f0(lad["R0"]["G"]), g["beta"]))
    ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws.row_dimensions[r].height = 30
    # firm tenants
    ws = wb.create_sheet("Firm tenants")
    _title(ws, "Firm tenants: Q1 2027 requests held to frozen budgets",
           "CS-3 3.2: the request, held to the whole cores the monthly budget pays for at the rate in force for the "
           "quarter (EUR %.2f per core-month from 1 January 2027). Budgets: January to March 2027 held at December "
           "2026 (freeze note of 27 November)." % M.rate_q1)
    _head(ws, 4, ["Tenant", "Cost centre", "Requested", "Monthly budget (EUR)", "Cores the budget pays for",
                  "Approved", "Held to budget"], [14, 30, 12, 14, 14, 12, 10])
    r = 5
    for t in M.firm:
        b = M.budget[t]["Jan 2027"]
        cap = int(math.floor(b / M.rate_q1 + 1e-9))
        _row(ws, r, [t, M.cc[t], M.requests[t], b, cap, c["approvals"][t],
                     "yes" if c["approvals"][t] < M.requests[t] else "no"],
             [None, None, "#,##0", "#,##0", "#,##0", "#,##0", None])
        r += 1
    _row(ws, r, ["Total", None, sum(M.requests.values()), None, None, g["F"], None],
         [None, None, "#,##0", None, None, "#,##0", None], bold=True)
    # storage
    ws = wb.create_sheet("cell-store")
    _title(ws, "cell-store: busiest-hour forecast from its drivers",
           "Least squares on every hour on file (%s hours, 2025-01-01 to 2026-12-13): cell-store's cores in use on "
           "the cell's machines in service and the batch tier's cores in use. Residual s.d. %.1f cores." % (
               f0(g["n"]), g["resid_sd"]))
    _head(ws, 4, ["Term", "Cores"], [40, 14])
    for k, (lab, v, nd) in enumerate([("Intercept", g["a"], 1), ("Per machine in service", g["c"], 4),
                                      ("Per batch core running", g["beta"], 5)]):
        _row(ws, 5 + k, [lab, round(v, nd)], [None, "0." + "0" * nd])
    _head(ws, 9, ["Quarter", "Busiest hour (UTC)", "cell-store at that hour", "Machines in service",
                  "batch-train cores", "Fitted"], [40, 18, 14, 14, 14, 12])
    r = 10
    for q in M.quarters:
        m = M.q == q
        k = int(max(range(int(m.sum())), key=lambda i: M.store[m][i]))
        tt = [t for t, mm in zip(M.t, m) if mm][k]
        fitted = g["a"] + g["c"] * float(M.mach[m][k]) + g["beta"] * float(M.batch[m][k])
        _row(ws, r, [q + (" (to 13 Dec)" if q == M.quarters[-1] else ""), tt.strftime("%Y-%m-%d %H:%M"),
                     float(M.store[m][k]), float(M.mach[m][k]), float(M.batch[m][k]), round(fitted, 1)],
             [None, None, "#,##0.0", "#,##0.0", "#,##0.0", "#,##0.0"])
        r += 1
    r += 1
    _row(ws, r, ["2027 Q1 at %s machines, batch at %s" % (f0(M.M1), f0(g["G"])), None, round(g["S"], 1)],
         [None, None, "#,##0.0"], bold=True)
    # Ask A
    ws = wb.create_sheet("Training lost 2026 Q3")
    _title(ws, "Batch training lost to preemptions, July to September 2026",
           "Core-hours each team's runs held from the checkpoint their resumption restored (its state as of the write "
           "start) to the preemption, resizes applied; share of all core-hours the team's runs held in the quarter.")
    _head(ws, 4, ["Team", "Core-hours held", "Core-hours on training lost to a preemption", "Share (%)"], [14, 16, 22, 11])
    r = 5
    for t in c["A"].teams:
        _row(ws, r, [t, K.round_to(c["held"][t], 100), K.round_to(c["redone"][t], 100), float(pct2(c["share"][t]))],
             [None, "#,##0", "#,##0", "0.00"])
        r += 1
    # Ask B
    ws = wb.create_sheet("Machine-hours lost 2026 Q3")
    _title(ws, "Machine-hours C-22 lost on machines that came back into service, July to September 2026",
           "Each removal to the asset's next registration in the cell, under whichever cluster id it came back with "
           "(hardware inventory). Machines retired after a fault are not included.")
    _head(ws, 4, ["Month", "Hardware faults", "Planned maintenance"], [14, 16, 20])
    for k, m in enumerate((7, 8, 9)):
        _row(ws, 5 + k, [MONTHS[m], K.round_to(c["lost"][("hw_fault", m)], 10),
                         K.round_to(c["lost"][("maintenance", m)], 10)], [None, "#,##0", "#,##0"])
    _row(ws, 8, ["Quarter", K.round_to(sum(c["lost"][("hw_fault", m)] for m in (7, 8, 9)), 10),
                 K.round_to(sum(c["lost"][("maintenance", m)] for m in (7, 8, 9)), 10)], [None, "#,##0", "#,##0"],
         bold=True)
    ws.cell(row=10, column=1, value="Machine-hours to the nearest 10; the quarter's row is summed before rounding.")
    ws.cell(row=10, column=1).font = Font(name="Calibri", italic=True, size=9, color="595959")
    # Ask C
    ws = wb.create_sheet("Chargeback 2026 Q3")
    _title(ws, "What each firm tenant owes for July to September 2026 (EUR)",
           "Chargeback policy 2026: each day, the cores held at 00:00 UTC (the review's reservation and every change "
           "the quota service had applied by then) at the rate in force that day, over the days in the month.")
    _head(ws, 4, ["Tenant", "Cost centre", "Q3 2026 (EUR, nearest 100)"], [14, 30, 18])
    r = 5
    for t in c["A"].firm:
        _row(ws, r, [t, M.cc[t], K.round_to(c["owe"][t], 100)], [None, None, "#,##0"])
        r += 1
    _row(ws, r, ["Eleven firm tenants", None, K.round_to(sum(c["owe"].values()), 100)], [None, None, "#,##0"], bold=True)
    ws.cell(row=r + 2, column=1, value="Rates: EUR 7.65 per core-month to 31 August, EUR 7.70 from 1 September 2026 "
            "(MAD1 reserved cores, rate card). Totals summed before rounding to the nearest EUR 100.")
    ws.cell(row=r + 2, column=1).font = Font(name="Calibri", italic=True, size=9, color="595959")
    # notes
    ws = wb.create_sheet("Notes")
    _title(ws, "Notes", "Compute Operations, MAD1. Prepared 14 December 2026 for the C-22 review of 15 December.")
    notes = [
        ("Sources", "CS-3 (rev. June 2026); capacity memo of 9 December; C-22 hardware plan rev. 11 December; capacity "
                    "requests for 2027 Q1 (window closed 4 December 18:00 Madrid); technology budgets with the freeze "
                    "note of 27 November; compute rate card; C-22 tenant directory; C-22 monitoring export and machines "
                    "series, 1 January 2025 to 13 December 2026 23:00 UTC, pulled 14 December 06:10 UTC."),
        ("2026 Q3 tables", "Batch scheduler quarter snapshot of 1 October 2026 (runs, checkpoints, resizes); MAD1 "
                           "cluster manager events for C-22 and the hardware inventory extract of 14 December; quota "
                           "tool changes and quota service reconciliation runs; chargeback policy 2026."),
        ("Cores", "Reservable cores throughout (CS-3 2.2). Reservations to the nearest 1,000 cores; the guarantee is the "
                  "largest whole thousand left (CS-3 5.1)."),
        ("Monitoring export", "17 March 2026 is in the export twice (the collector re-sent the day) and was counted "
                              "once; 19 August 2026 01:00 to 07:59 UTC has no rows (collector outage)."),
        ("Rounding", "Core-hours to the nearest 100, shares to two decimals, machine-hours to the nearest 10, charges to "
                     "the nearest EUR 100; every figure rounded once, from unrounded working."),
    ]
    _head(ws, 4, ["Item", "Note"], [20, 110])
    for k, (a, b) in enumerate(notes):
        _row(ws, 5 + k, [a, b])
        ws.cell(row=5 + k, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[5 + k].height = 44 if len(b) > 200 else (30 if len(b) > 100 else 15)
    for w in wb.worksheets:
        w.page_setup.orientation = "landscape"
        w.page_setup.paperSize = w.PAPERSIZE_A4
        w.sheet_properties.pageSetUpPr.fitToPage = True
        w.page_setup.fitToWidth = 1
        w.page_setup.fitToHeight = 0
        if w.title not in ("2027 Q1", "Notes"):
            w.freeze_panes = "A5"
    wb.save(path)
    PP.normalize_ooxml(path, K.PEOPLE["requester"], K.PEOPLE["requester"], "2026-12-14T08:31:07Z",
                       "2026-12-14T17:48:52Z", excel=True, company=K.COMPANY)


def write_deck(path, c, figs):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
    from pptx.util import Emu, Inches, Pt
    g, M = c["g"], c["M"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    ink, muted = RGBColor(0x1D, 0x23, 0x30), RGBColor(0x5A, 0x62, 0x70)

    def text(x, y, w, h, t, size, bold=False, color=ink):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = t
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.font.name = "Calibri"
        return tb
    F, P_, G = K.round_to(g["F"], 1000), K.round_to(g["platform"], 1000), g["G"]
    text(0.6, 0.35, 12.2, 0.8, "C-22 in 2027 Q1: %s cores guaranteed to batch training" % f0(G), 28, bold=True)
    text(0.6, 1.1, 12.2, 0.6, "How the quarter's %s reservable cores divide (%s machines x 96 x 1.45), to the nearest "
         "1,000 cores" % (f0(M.R1), f0(M.M1)), 15, color=muted)
    cd = CategoryChartData()
    cd.categories = ["2027 Q1"]
    parts = [("Firm tenants", F), ("Platform services (cell-store %s, cell-ctl %s)" % (
        f0(K.round_to(g["S"], 1000)), f0(K.round_to(g["ctl"], 1000))), P_), ("Batch training guarantee", G)]
    for lab, v in parts:
        cd.add_series(lab, (v / 1000.0,))
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(0.6), Inches(1.9), Inches(12.2), Inches(3.4), cd)
    ch = gf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(12)
    va = ch.value_axis
    va.has_title = True
    va.axis_title.text_frame.text = "Reservable cores (thousands)"
    va.axis_title.text_frame.paragraphs[0].font.size = Pt(12)
    va.maximum_scale = math.ceil(M.R1 / 1000.0 / 50.0) * 50
    va.minimum_scale = 0
    va.tick_labels.font.size = Pt(11)
    ch.category_axis.tick_labels.font.size = Pt(12)
    colours = [RGBColor(0x8C, 0x9B, 0xB0), RGBColor(0xE0, 0x8E, 0x2B), RGBColor(0x1F, 0x4E, 0x79)]
    plot = ch.plots[0]
    plot.gap_width = 40
    for ser, col, (lab, v) in zip(plot.series, colours, parts):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = col
        ser.data_labels.show_value = True
        ser.data_labels.number_format = '#,##0"k"'
        ser.data_labels.number_format_is_linked = False
        ser.data_labels.font.size = Pt(12)
        ser.data_labels.font.bold = True
        ser.data_labels.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    rows = [("Firm tenants", F, "eleven requests held to frozen budgets at EUR %.2f; three held below their requests" %
             M.rate_q1),
            ("Platform services", P_, "cell-store %s, which rises with the batch tier it serves, and cell-ctl %s" % (
                f0(K.round_to(g["S"], 1000)), f0(K.round_to(g["ctl"], 1000)))),
            ("Batch training guarantee", G, "largest whole thousand left once cell-store carries the batch tier's "
             "load"),
            ("Not reserved", round(g["unreserved"]), "")]
    tb = s.shapes.add_table(len(rows) + 1, 3, Inches(0.6), Inches(5.45), Inches(12.2), Inches(1.6)).table
    tb.columns[0].width, tb.columns[1].width, tb.columns[2].width = Inches(3.0), Inches(1.6), Inches(7.6)
    for k, h in enumerate(("Part", "Cores", "How it is set")):
        tb.cell(0, k).text = h
    for i, (lab, v, note) in enumerate(rows, 1):
        tb.cell(i, 0).text = lab
        tb.cell(i, 1).text = f0(v)
        tb.cell(i, 2).text = note
    for i in range(len(rows) + 1):
        for k in range(3):
            for p in tb.cell(i, k).text_frame.paragraphs:
                p.font.size = Pt(11)
                p.font.name = "Calibri"
    text(0.6, 7.08, 12.2, 0.3, "Sources: C-22 hardware plan rev. 11 December 2026; 2027 Q1 requests and budgets as "
         "frozen 27 November; MAD1 rate card from 1 January 2027; C-22 monitoring export to 13 December 2026 23:00 UTC. "
         "Compute Operations, MAD1, for the review of 15 December 2026 (CS-3).", 9, color=muted)
    prs.core_properties.author = K.PEOPLE["requester"]
    prs.core_properties.title = "C-22 capacity review 2027 Q1"
    prs.save(path)
    _fix_embedded(path, "2026-12-14T15:02:40Z")
    PP.normalize_ooxml(path, K.PEOPLE["requester"], K.PEOPLE["requester"], "2026-12-14T15:02:40Z",
                       "2026-12-14T17:55:13Z", title="C-22 capacity review 2027 Q1", company=K.COMPANY)


def _fix_xlsx(data, stamp):
    import io
    import re
    import zipfile
    zin = zipfile.ZipFile(io.BytesIO(data))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for i in zin.infolist():
            d = zin.read(i.filename)
            if i.filename == "docProps/core.xml":
                for tag in (b"created", b"modified"):
                    d = re.sub(rb"(<dcterms:%s[^>]*>)[^<]*(</dcterms:%s>)" % (tag, tag),
                               rb"\g<1>" + stamp.encode() + rb"\g<2>", d)
            zi = zipfile.ZipInfo(i.filename, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, d)
    return buf.getvalue()


def _fix_embedded(path, stamp):
    """The chart's embedded workbook carries the time it was written; set it to the deck's own time."""
    import zipfile
    zin = zipfile.ZipFile(path)
    items = [(i.filename, zin.read(i.filename)) for i in zin.infolist()]
    zin.close()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, d in items:
            if n.startswith("ppt/embeddings/") and n.endswith(".xlsx"):
                d = _fix_xlsx(d, stamp)
            if n.startswith("ppt/slideMasters/") and n.endswith(".xml"):
                d = d.replace('char="\u2013"'.encode(), 'char="\u2022"'.encode())   # the template's bullet
            zi = zipfile.ZipInfo(n, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, d)


def main(target, golden, figures_path):
    os.makedirs(golden, exist_ok=True)
    c = compute(target)
    figs = figures(c)
    write_book(os.path.join(golden, K.GOLD_BOOK), c, figs)
    write_deck(os.path.join(golden, K.GOLD_DECK), c, figs)
    with open(figures_path, "w") as fh:
        json.dump(figs, fh, indent=2, sort_keys=False)
        fh.write("\n")
    return c, figs
