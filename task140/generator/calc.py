"""Every graded figure and every rung, computed only from the shipped pack (the folder target/): the 2027 Q1 batch
guarantee and the reservations it is struck after, the ladder's wrong routes, and the three 2026 Q3 asks with their
stops. deliver.py writes the goldens from these functions; checks.py asserts on them; verify.py recomputes the graded
figures on its own code path."""
from __future__ import annotations

import csv
import datetime as dt
import json
import math
import os
import re

import numpy as np
import openpyxl

import common as K

UTC = dt.timezone.utc


def P(target, name):
    return os.path.join(target, name)


def rows(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def tparse(s):
    return dt.datetime.strptime(s, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=UTC)


def floor1000(x):
    return int(math.floor(x / 1000.0 + 1e-12)) * 1000


def near1000(x):
    return K.round_to(x, 1000)


def quarter_of(t):
    return "%dQ%d" % (t.year, (t.month - 1) // 3 + 1)


# ---------------------------------------------------------------------------------------------------- loading the main call's inputs

class Pack:
    pass


def docx_text(path):
    import docx
    d = docx.Document(path)
    return "\n".join(p.text for p in d.paragraphs)


def load(target, dedupe=True, gap="drop"):
    """The main call's inputs, as a careful reader takes them from the pack."""
    M = Pack()
    d = json.load(open(P(target, K.DIRECTORY), encoding="utf-8"))
    M.classes = {t["id"]: t["class"] for t in d["tenants"]}
    M.firm = [t["id"] for t in d["tenants"] if t["class"] == "firm"]
    M.platform = [t["id"] for t in d["tenants"] if t["class"] == "platform"]
    M.cc = {t["id"]: t["cost_centre"] for t in d["tenants"] if t["class"] == "firm"}
    memo = docx_text(P(target, K.MEMO))
    M.overcommit = float(re.search(r"overcommit for 2027 Q1 is ([0-9.]+)", memo).group(1))
    wb = openpyxl.load_workbook(P(target, K.HWPLAN), data_only=True)
    ws = wb["By quarter"]
    M.machines_q = {}
    for r in ws.iter_rows(values_only=True):
        if r and isinstance(r[0], str) and re.fullmatch(r"\d{4}Q\d", r[0]):
            M.machines_q[r[0]] = int(r[3])
    M.M1 = M.machines_q[K.NEXT_Q]
    M.racks = []
    for r in wb["Racks"].iter_rows(values_only=True):
        if r and isinstance(r[0], str) and re.fullmatch(r"2[A-M]-\d\d", r[0]):
            M.racks.append((r[0], r[4], int(r[5])))
    M.R1 = int(round(M.M1 * K.CORES * M.overcommit))
    M.requests = {}
    for r in rows(P(target, K.REQUESTS)):
        if r["quarter"] == K.NEXT_Q:
            assert r["tenant"] not in M.requests
            M.requests[r["tenant"]] = int(r["cores_requested"])
    wb = openpyxl.load_workbook(P(target, K.BUDGETS), data_only=True)
    ws = wb["Compute budgets"]
    head = None
    M.budget_lines = []
    for r in ws.iter_rows(values_only=True):
        if r and r[0] == "CC":
            head = list(r)
            continue
        if head and r and r[0]:
            M.budget_lines.append(dict(zip(head, r)))
    M.budget = {}
    M.budget_all = {}
    for t in M.firm:
        code = M.cc[t].split()[0].replace("CC-", "")
        line = [x for x in M.budget_lines if str(x["CC"]) == code and x["Budget line"] == "MAD1 C-22 reserved cores"]
        assert len(line) == 1, (t, line)
        M.budget[t] = {m: line[0][m] for m in ("Jan 2027", "Feb 2027", "Mar 2027")}
        M.budget_all[t] = line[0]
    M.rates = [(dt.date.fromisoformat(r["effective_from"]), float(r["eur_per_unit_month"]))
               for r in rows(P(target, K.RATECARD)) if r["resource"] == "reserved core" and r["site"] == "MAD1"]
    M.rates.sort()
    M.rate_q1 = rate_on(M.rates, dt.date(2027, 1, 1))
    # the monitoring export and the machines series
    seen = set()
    use = {}
    with open(P(target, K.USAGE), encoding="utf-8") as fh:
        next(fh)
        for line in fh:
            if dedupe and line in seen:
                continue
            seen.add(line)
            h, _, t, v = line.rstrip("\n").split(",")
            use.setdefault(t, {})
            if not dedupe and h in use[t]:
                use[t][h + "#2"] = float(v)
            else:
                use[t][h] = float(v)
    cellh = {}
    for r in rows(P(target, K.CELLHOURLY)):
        cellh[r["hour_start_utc"]] = (float(r["machines_installed"]), float(r["machines_in_service"]))
    hrs = sorted(k for k in use["cell-store"])
    if gap == "interpolate":
        hrs, use, cellh = _interpolate_gap(hrs, use, cellh)
    M.hours = hrs
    M.t = [tparse(h[:20]) for h in hrs]
    M.q = np.array([quarter_of(x) for x in M.t])
    M.store = np.array([use["cell-store"][h] for h in hrs])
    M.ctl = np.array([use["cell-ctl"][h] for h in hrs])
    M.batch = np.array([use["batch-train"][h] for h in hrs])
    M.mach = np.array([cellh[h[:20]][1] for h in hrs])
    M.inst = np.array([cellh[h[:20]][0] for h in hrs])
    M.use = use
    M.ledger = {}
    for r in rows(P(target, K.LEDGER)):
        M.ledger[(r["quarter"], r["tenant"])] = int(r["reserved_cores"])
    M.quarters = sorted({q for q, _ in M.ledger})
    return M


def _interpolate_gap(hrs, use, cellh):
    """Fill the collector outage's seven hours by straight lines between the hours either side."""
    t = [tparse(h) for h in hrs]
    out_h = []
    for a, b, ha, hb in zip(t[:-1], t[1:], hrs[:-1], hrs[1:]):
        out_h.append(ha)
        n = int((b - a).total_seconds() // 3600)
        for k in range(1, n):
            hk = (a + dt.timedelta(hours=k)).strftime("%Y-%m-%dT%H:%M:%SZ")
            w = k / n
            for s in use:
                if ha in use[s] and hb in use[s]:
                    use[s][hk] = use[s][ha] * (1 - w) + use[s][hb] * w
            cellh[hk] = tuple(cellh[ha][i] * (1 - w) + cellh[hb][i] * w for i in range(2))
            out_h.append(hk)
    out_h.append(hrs[-1])
    return out_h, use, cellh


def rate_on(rates, d):
    v = None
    for e, x in rates:
        if e <= d:
            v = x
    return v


# ---------------------------------------------------------------------------------------------------- the main call

def firm_approvals(M, rate=None, budget_month="Jan 2027", hold=True, rounding="floor"):
    rate = rate or M.rate_q1
    out = {}
    for t in M.firm:
        b = M.budget_all[t][budget_month] if budget_month in M.budget_all[t] else None
        if budget_month == "avg2026":
            b = sum(M.budget_all[t][m] for m in M.budget_all[t] if isinstance(m, str) and m.endswith("2026")) / 12.0
        x = b / rate
        cap = int(math.floor(x + 1e-9)) if rounding == "floor" else int(math.floor(x + 0.5))
        out[t] = min(M.requests[t], cap) if hold else M.requests[t]
    return out


def qmax(M, series, q):
    v = series[M.q == q]
    return float(v.max())


def line_forecast(vals):
    x = np.arange(len(vals))
    a, b = np.polyfit(x, vals, 1)
    return float(a * len(vals) + b)


def fit(M, cols, mask=None, y=None):
    y = M.store if y is None else y
    m = np.ones(len(y), bool) if mask is None else mask
    X = np.column_stack([np.ones(m.sum())] + [c[m] for c in cols])
    co, *_ = np.linalg.lstsq(X, y[m], rcond=None)
    res = y[m] - X @ co
    return co, res


def ctl_forecast(M, how="trend4"):
    last4 = [qmax(M, M.ctl, q) for q in M.quarters[-4:]]
    if how == "trend4":
        return line_forecast(last4)
    if how == "last":
        return last4[-1]
    if how == "mean4":
        return float(np.mean(last4))
    if how == "machines":
        mx = [float(M.mach[M.q == q][np.argmax(M.ctl[M.q == q])]) for q in M.quarters]
        co = np.polyfit(mx, [qmax(M, M.ctl, q) for q in M.quarters], 1)
        return float(np.polyval(co, M.M1))
    raise ValueError(how)


def golden(M, mask=None, allowance=0.0, machines=None, use_installed=False, ctl_how="trend4", F=None, R=None,
           G_in_store="guarantee"):
    """The guarantee: cell-store's use specified on machines in service and batch cores running, read at the quarter's
    machines with the batch tier at its own guarantee, solved together with the guarantee."""
    mach = M.inst if use_installed else M.mach
    co, res = fit(M, [mach, M.batch], mask)
    a, c, beta = (float(x) for x in co)
    Mq = M.M1 if machines is None else machines
    F = sum(firm_approvals(M).values()) if F is None else F
    R = M.R1 if R is None else R
    ctl = ctl_forecast(M, ctl_how)
    Gu = (R - F - ctl - a - c * Mq - allowance) / (1 + beta)
    G = floor1000(Gu)
    Gs = G if G_in_store == "guarantee" else Gu
    S = a + c * Mq + beta * Gs + allowance
    return dict(a=a, c=c, beta=beta, resid_sd=float(res.std()), resid_max=float(res.max()), n=int(len(res)),
                R=R, F=F, ctl=ctl, Gu=Gu, G=G, S=S, platform=S + ctl, unreserved=R - F - ctl - S - G, M1=Mq)


def from_terms(M, co, F=None, ctl=None):
    """The golden's solve on another fit's terms (intercept, per machine, per batch core)."""
    a, c, beta = (float(x) for x in co)
    F = sum(firm_approvals(M).values()) if F is None else F
    ctl = ctl_forecast(M) if ctl is None else ctl
    Gu = (M.R1 - F - ctl - a - c * M.M1) / (1 + beta)
    G = floor1000(Gu)
    S = a + c * M.M1 + beta * G
    return dict(a=a, c=c, beta=beta, F=F, ctl=ctl, Gu=Gu, G=G, S=S, platform=S + ctl, unreserved=M.R1 - F - ctl - S - G)


def other_fits(M):
    """Fits a careful analyst might prefer once the residual's low hours are seen: the Tuesday restart window left
    out, the hours more than three residual s.d. below the line left out, a Huber fit, and hour-of-week terms."""
    tt = [tparse(h[:20]) for h in M.hours]
    uw = np.array([x.weekday() for x in tt])
    uh = np.array([x.hour for x in tt])
    out = {"restart window left out": golden(M, mask=~((uw == 1) & (uh >= 3) & (uh < 5)))}
    co, res = fit(M, [M.mach, M.batch])
    out["low hours left out"] = golden(M, mask=res > -3 * res.std())
    X = np.column_stack([np.ones(len(M.store)), M.mach, M.batch])
    b = co.copy()
    for _ in range(50):
        r = M.store - X @ b
        k = 1.345 * 1.4826 * float(np.median(np.abs(r - np.median(r))))
        w = np.where(np.abs(r) <= k, 1.0, k / np.maximum(np.abs(r), 1e-9))
        b = np.linalg.solve(X.T @ (X * w[:, None]), (X * w[:, None]).T @ M.store)
    out["Huber"] = from_terms(M, b)
    how = uw * 24 + uh
    D = np.column_stack([(how == k).astype(float) for k in range(1, 168)])
    bd, *_ = np.linalg.lstsq(np.column_stack([X, D]), M.store, rcond=None)
    out["hour-of-week terms"] = from_terms(M, [bd[0] + float(np.mean(D @ bd[3:])), bd[1], bd[2]])
    return out


def guarantee_given_store(M, S, F=None, ctl=None):
    F = sum(firm_approvals(M).values()) if F is None else F
    ctl = ctl_forecast(M) if ctl is None else ctl
    Gu = M.R1 - F - ctl - S
    return dict(Gu=Gu, G=floor1000(Gu), S=S, F=F, ctl=ctl)


def ladder(M):
    """Every route the design prices, each a complete pipeline on CS-3's estimand."""
    F = sum(firm_approvals(M).values())
    Fw = sum(firm_approvals(M, hold=False).values())
    ctl = ctl_forecast(M)
    last4 = [qmax(M, M.store, q) for q in M.quarters[-4:]]
    all8 = [qmax(M, M.store, q) for q in M.quarters]
    out = {}
    out["R-1"] = guarantee_given_store(M, line_forecast(last4), F=Fw)
    out["R0"] = guarantee_given_store(M, line_forecast(last4))
    out["R0b"] = guarantee_given_store(M, last4[-1])
    out["R0c"] = guarantee_given_store(M, line_forecast(all8))
    out["R0d"] = guarantee_given_store(M, float(np.mean(last4)))
    full_q = [q for q in M.quarters if q != M.quarters[-1]]
    out["R0e"] = guarantee_given_store(M, line_forecast([qmax(M, M.store, q) for q in full_q[-4:]]))
    q4 = M.quarters[-1]
    m4 = M.q == q4
    hb = int(np.argmax(M.store[m4]))
    out["R1"] = guarantee_given_store(M, float(M.store[m4][hb]) / float(M.mach[m4][hb]) * M.M1)
    R_q4 = int(round(M.machines_q[q4] * K.CORES * M.overcommit))
    out["R1b"] = guarantee_given_store(M, last4[-1] / R_q4 * M.R1)
    out["R1c"] = guarantee_given_store(M, M.ledger[(q4, "cell-store")] / M.machines_q[q4] * M.M1)
    g = golden(M)
    a, c, beta = g["a"], g["c"], g["beta"]
    out["R2b"] = guarantee_given_store(M, a + c * M.M1 + beta * out["R0"]["Gu"])
    out["R2e"] = guarantee_given_store(M, a + c * M.M1 + beta * M.ledger[(q4, "batch-train")])
    m_now = sum(n for _, d, n in M.racks if d is not None and _d(d) < dt.date(2026, 12, 14))
    Gu = (M.R1 - F - ctl - a - c * m_now) / (1 + beta)
    out["R2d"] = dict(Gu=Gu, G=floor1000(Gu), S=a + c * m_now + beta * floor1000(Gu))
    co, _ = fit(M, [M.batch])
    Gu = (M.R1 - F - ctl - co[0]) / (1 + co[1])
    out["R2"] = dict(Gu=Gu, G=floor1000(Gu), S=co[0] + co[1] * floor1000(Gu), coef=[float(x) for x in co])
    lvl = float(np.mean(M.store[m4] - beta * M.batch[m4]))
    Gu = (M.R1 - F - ctl - lvl) / (1 + beta)
    out["R2c"] = dict(Gu=Gu, G=floor1000(Gu), S=lvl + beta * floor1000(Gu))
    co, _ = fit(M, [M.mach])
    out["Monly"] = guarantee_given_store(M, co[0] + co[1] * M.M1)
    ratio = last4[-1] / M.ledger[(q4, "batch-train")]
    Gu = (M.R1 - F - ctl) / (1 + ratio)
    out["O1"] = dict(Gu=Gu, G=floor1000(Gu), S=ratio * floor1000(Gu))
    # the ledger's past cell-store reservations against the machines at each quarter's end and the guarantee
    end_m = []
    for q in M.quarters:
        nq = _next_q(q)
        end_m.append(M.machines_q[nq] if nq != K.NEXT_Q else
                     sum(n for _, d, n in M.racks if d is not None and _d(d) < dt.date(2026, 12, 14)))
    Xl = np.array([[1.0, m, M.ledger[(q, "batch-train")]] for q, m in zip(M.quarters, end_m)])
    cl, *_ = np.linalg.lstsq(Xl, np.array([M.ledger[(q, "cell-store")] for q in M.quarters], float), rcond=None)
    Gu = (M.R1 - F - ctl - cl[0] - cl[1] * M.M1) / (1 + cl[2])
    out["LedgerFit"] = dict(Gu=Gu, G=floor1000(Gu), coef=[float(x) for x in cl])
    Mq, Bq, Sq = [], [], []
    for q in M.quarters:
        mm = M.q == q
        k = int(np.argmax(M.store[mm]))
        Mq.append(float(M.mach[mm][k]))
        Bq.append(float(M.batch[mm][k]))
        Sq.append(float(M.store[mm][k]))
    cq, *_ = np.linalg.lstsq(np.column_stack([np.ones(len(Sq)), Mq, Bq]), np.array(Sq), rcond=None)
    Gu = (M.R1 - F - ctl - cq[0] - cq[1] * M.M1) / (1 + cq[2])
    out["QmaxFit"] = dict(Gu=Gu, G=floor1000(Gu), coef=[float(x) for x in cq])
    out["R3"] = g
    return out


def _d(x):
    return x.date() if isinstance(x, dt.datetime) else x


def _next_q(q):
    y, n = int(q[:4]), int(q[-1])
    return "%dQ%d" % (y + (n == 4), 1 if n == 4 else n + 1)


# ---------------------------------------------------------------------------------------------------- Ask A: training lost to preemptions

Q3S = dt.datetime(2026, 7, 1, tzinfo=UTC)
Q3E = dt.datetime(2026, 10, 1, tzinfo=UTC)


class Asks:
    pass


def load_asks(target):
    A = Asks()
    d = json.load(open(P(target, K.DIRECTORY), encoding="utf-8"))
    A.teams = [x["id"] for t in d["tenants"] if t["class"] == "batch" for x in t["teams"]]
    A.firm = [t["id"] for t in d["tenants"] if t["class"] == "firm"]
    A.runs = rows(P(target, K.RUNS))
    A.ckpt = {r["checkpoint_id"]: r for r in rows(P(target, K.CHECKPOINTS))}
    A.ckpt_by_run = {}
    for r in A.ckpt.values():
        A.ckpt_by_run.setdefault(r["run_id"], []).append(r)
    A.resizes = {}
    for r in rows(P(target, K.RESIZES)):
        A.resizes.setdefault(r["run_id"], []).append((tparse(r["at_utc"]), int(r["cores_after"])))
    for v in A.resizes.values():
        v.sort()
    A.mev = rows(P(target, K.MEVENTS))
    wb = openpyxl.load_workbook(P(target, K.INVENTORY), data_only=True)
    A.asset_of = {}
    for r in wb["Cluster ids"].iter_rows(values_only=True):
        if r and isinstance(r[0], str) and r[0].startswith("SRV-"):
            A.asset_of[int(r[1])] = r[0]
    A.status = {}
    for r in wb["Assets"].iter_rows(values_only=True):
        if r and isinstance(r[0], str) and r[0].startswith("SRV-"):
            A.status[r[0]] = r[7]
    A.ledger = {}
    for r in rows(P(target, K.LEDGER)):
        A.ledger[(r["quarter"], r["tenant"])] = int(r["reserved_cores"])
    A.changes = rows(P(target, K.QCHANGES))
    A.qruns = rows(P(target, K.QRUNS))
    A.rates = sorted((dt.date.fromisoformat(r["effective_from"]), float(r["eur_per_unit_month"]))
                     for r in rows(P(target, K.RATECARD)) if r["resource"] == "reserved core" and r["site"] == "MAD1")
    return A


def segments(A, run, resizes=True):
    t0 = tparse(run["started_utc"])
    t1 = tparse(run["ended_utc"]) if run["ended_utc"] else Q3E
    pts = [(t0, int(run["cores"]))] + (A.resizes.get(run["run_id"], []) if resizes else [])
    out = []
    for k, (a, v) in enumerate(pts):
        b = pts[k + 1][0] if k + 1 < len(pts) else t1
        out.append((a, min(b, t1), v))
    return out


def core_hours(segs, a, b):
    tot = 0.0
    for s, e, v in segs:
        lo, hi = max(s, a), min(e, b)
        if hi > lo:
            tot += v * (hi - lo).total_seconds() / 3600.0
    return tot


def ask_a(A, mode="golden"):
    """mode: golden (resizes, the restore point the resumption took, write start); s1 (cores at start, the preempted
    run's own newest committed checkpoint); s2 (resizes, own newest committed); s2b (resizes, own newest write,
    committed or not); s3 (restore point followed, its commit time as the training state)."""
    succ = {r["resumed_from"]: r for r in A.runs if r["resumed_from"]}
    held = {t: 0.0 for t in A.teams}
    redone = {t: 0.0 for t in A.teams}
    for r in A.runs:
        segs = segments(A, r, resizes=(mode != "s1"))
        held[r["team"]] += core_hours(segs, Q3S, Q3E)
        if r["end_state"] != "PREEMPTED":
            continue
        end = tparse(r["ended_utc"])
        if not (Q3S <= end < Q3E):
            continue
        start = tparse(r["started_utc"])
        own = sorted(A.ckpt_by_run.get(r["run_id"], []), key=lambda c: c["write_started_utc"])
        if mode in ("golden", "s3"):
            c = A.ckpt[succ[r["run_id"]]["restore_checkpoint"]]
            if c["run_id"] == r["run_id"]:
                t0 = tparse(c["write_started_utc"] if mode == "golden" else c["committed_utc"])
            else:
                t0 = start
        elif mode in ("s1", "s2"):
            com = [c for c in own if c["committed_utc"] and tparse(c["committed_utc"]) <= end]
            t0 = tparse(com[-1]["write_started_utc"]) if com else start
        elif mode == "s2b":
            t0 = tparse(own[-1]["write_started_utc"]) if own else start
        elif mode == "redo":
            # the resumed run's core-hours spent getting back to where the preempted run stopped
            c = A.ckpt[succ[r["run_id"]]["restore_checkpoint"]]
            t0 = tparse(c["write_started_utc"]) if c["run_id"] == r["run_id"] else start
            s = succ[r["run_id"]]
            s0 = tparse(s["started_utc"])
            redone[r["team"]] += core_hours(segments(A, s), s0, min(Q3E, s0 + (end - max(t0, Q3S))))
            continue
        else:
            raise ValueError(mode)
        redone[r["team"]] += core_hours(segs, max(t0, Q3S), Q3E)
    share = {t: 100.0 * redone[t] / held[t] for t in A.teams}
    return redone, share, held


# ---------------------------------------------------------------------------------------------------- Ask B: machine-hours lost

def ask_b(A, mode="golden", dedupe=True):
    """mode: golden (each removal to its asset's next registration under any id, from the inventory's id history);
    s1 (matched on machine id); s2 (matched on machine id, and an unmatched removal paired with the next registration
    in its rack that no same-id return claims)."""
    ev = []
    seen = set()
    for r in A.mev:
        key = tuple(r.values())
        if dedupe and key in seen:
            continue
        seen.add(key)
        ev.append(dict(r, t=tparse(r["time_utc"]), mid=int(r["machine_id"])))
    ev.sort(key=lambda x: x["t"])
    regs = [x for x in ev if x["event"] == "REGISTERED"]
    out = {(c, m): 0.0 for c in ("hw_fault", "maintenance") for m in (7, 8, 9)}
    removed = [x for x in ev if x["event"] == "REMOVED"]
    back = {}
    if mode == "golden":
        for x in removed:
            asset = A.asset_of[x["mid"]]
            back[id(x)] = next((y for y in regs if y["t"] > x["t"] and A.asset_of[y["mid"]] == asset), None)
    else:
        claimed = set()
        for x in removed:
            y = next((y for y in regs if y["t"] > x["t"] and y["mid"] == x["mid"]), None)
            back[id(x)] = y
            if y is not None:
                claimed.add(id(y))
        if mode == "s2":
            for x in removed:
                if back[id(x)] is None:
                    y = next((y for y in regs if y["t"] > x["t"] and y["rack"] == x["rack"] and id(y) not in claimed),
                             None)
                    back[id(x)] = y
                    if y is not None:
                        claimed.add(id(y))
    for x in removed:
        y = back[id(x)]
        if y is not None:
            out[(x["cause"], x["t"].month)] += (y["t"] - x["t"]).total_seconds() / 3600.0
    return out


# ---------------------------------------------------------------------------------------------------- Ask C: Q3 charges

def ask_c(A, mode="golden"):
    """mode: golden (a change counts from the reconciliation run that applied it; the rate in force each day);
    s1 (from the request time); s2 (application time, Q3's opening rate throughout)."""
    applied = {}
    for r in A.qruns:
        for cid in filter(None, r["changes_applied"].split(";")):
            applied[cid] = tparse(r["run_started_utc"])
    out = {}
    for t in A.firm:
        base = A.ledger[("2026Q3", t)]
        ch = [(applied[c["change_id"]] if mode != "s1" else tparse(c["requested_utc"]), int(c["delta_cores"]))
              for c in A.changes if c["tenant"] == t]
        tot = 0.0
        d = Q3S
        while d < Q3E:
            held = base + sum(v for when, v in ch if when <= d)
            rate = rate_on(A.rates, d.date()) if mode != "s2" else rate_on(A.rates, Q3S.date())
            nxt = dt.date(d.year + (d.month == 12), 1 if d.month == 12 else d.month + 1, 1)
            dim = (nxt - dt.date(d.year, d.month, 1)).days
            tot += held * rate / dim
            d += dt.timedelta(days=1)
        out[t] = tot
    return out
