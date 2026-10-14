"""Write the bundle's data files from the world: the three feeds in their systems' own layouts, the lookups, the
dictionary, the manifest and the two files the solution never reads."""
import csv
import datetime as dt
import json

import numpy as np

import feeds as FD
import latent as LT
import roster as RO
import world as W

FILES = {
    "compact": "regional_icu_surge_compact_2025.pdf",
    "annex": "compact_data_annex_2025-11.pdf",
    "regions": "compact_regions_counties.csv",
    "crosswalk": "facility_crosswalk.csv",
    "aliases": "transfer_center_name_aliases.csv",
    "codes": "transfer_center_reason_codes.csv",
    "survey_old": "capacity_survey_2025-11.csv",
    "survey_new": "capacity_survey_2025-12-01_2026-01-19.csv",
    "board": "bed_status_events_2025-11-01_2026-01-20.jsonl",
    "tc_old": "transfer_center_calls_2025-11-01_2026-01-04.xlsx",
    "tc_new": "transfer_center_calls_2026-01-05_2026-01-19.csv",
    "dict": "data_dictionary.xlsx",
    "manifest": "export_manifest_2026-01-20.csv",
    "federal": "federal_hospital_respiratory_weekly_2025-12_2026-01.csv",
    "history": "compact_activations_2024-25.csv",
}
CENTER_LEVEL_OLD = {"Adult ICU": "Adult ICU", "Step-down": "Step-down/PCU", "Pediatric ICU": "Pediatric ICU"}
CENTER_LEVEL_NEW = {"Adult ICU": "ADULT_ICU", "Step-down": "STEP_DOWN", "Pediatric ICU": "PEDS_ICU"}


def _csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, lineterminator="\r\n")
        w.writerow(header)
        w.writerows(rows)
    return len(rows)


# ------------------------------------------------------------------------------------------------- lookups
def write_regions(path):
    rows = []
    for code, name, *_ in W.REGIONS:
        for c in sorted(W.COUNTIES[code] + W.REGION_FILE_EXTRA.get(code, [])):
            rows.append((c, "47" + W.FIPS[c], code, name))
    rows.sort(key=lambda r: r[1])
    return _csv(path, ["county", "county_fips", "region_code", "region_name"], rows)


def write_crosswalk(hosp, path):
    rng = np.random.default_rng(W.SEED + 50)
    rows = []
    for h in sorted(hosp, key=lambda h: h["licence"]):
        rows.append((h["licence"], h["name"], h["county"], h["fips"], "America/New_York" if h["tz"] == "Eastern" else
                     "America/Chicago", h["board_code"], h["board_code_old"], h["center_code"],
                     h["beds"] + int(rng.integers(0, 7)), int(rng.choice([2019, 2019, 2020, 2021, 2022, 2023]))))
    return _csv(path, ["licence_no", "facility_name", "county", "county_fips", "time_zone", "board_code",
                       "board_code_previous", "center_code", "adult_icu_beds_licensed", "member_since"], rows)


def write_aliases(hosp, path):
    tzname = {"Eastern": "America/New_York", "Central": "America/Chicago"}
    rows = [(nm, fid, next(h["name"] for h in hosp if h["licence"] == fid), cty, "TN",
             tzname[next(h["tz"] for h in hosp if h["licence"] == fid)], "Y") for nm, fid, cty in RO.aliases(hosp)]
    rng = np.random.default_rng(W.SEED + 51)
    for f in RO.nonmembers():
        names = {f["name"], f["name"].upper()}
        if rng.random() < 0.6:
            names.add(f["name"].replace("Community Hospital", "Comm Hosp").replace("Medical Center", "Med Ctr")
                      .replace("Hospital", "Hosp"))
        for nm in sorted(names):
            rows.append((nm, f["fid"], f["name"], f["county"], f["state"], tzname[f["tz"]], "N"))
    rows.sort(key=lambda r: (r[1], r[0]))
    return _csv(path, ["typed_name", "facility_id", "facility_name", "county", "state", "time_zone",
                       "compact_member"], rows)


def write_codes(path):
    rows = [("NB", "Declined: no bed. No staffed bed of the level of care requested.", "2023-07-01", "2025-11-30"),
            ("NS", "Declined: no staff. A bed but no nurse to staff it.", "2023-07-01", "2025-11-30"),
            ("SP", "Declined: specialty. The service the patient needs is not offered.", "2023-07-01", "2025-11-30"),
            ("OT", "Declined: other reason.", "2023-07-01", "2025-11-30"),
            ("CAP", "Declined: capacity. No staffed bed of the level of care requested, for want of a bed or of "
                    "staff.", "2025-12-01", ""),
            ("SVC", "Declined: service. The service the patient needs is not offered.", "2025-12-01", ""),
            ("PT", "Declined: patient. A reason specific to the patient (insurance, family, condition).",
             "2025-12-01", ""),
            ("OTH", "Declined: other reason.", "2025-12-01", "")]
    return _csv(path, ["code", "meaning", "in_use_from", "in_use_to"], rows)


# ------------------------------------------------------------------------------------------------- survey
def write_survey(hosp, S, path_old, path_new):
    by = {h["licence"]: h for h in hosp}
    old, new = [], []
    sid = 7100000
    ids = np.random.default_rng(W.SEED + 22)          # the survey numbers every hospital's submissions in one series
    for r in sorted(S, key=lambda r: (r["date"], r["submitted"], r["licence"])):
        h = by[r["licence"]]
        sub = r["submitted"]
        if r["date"] < W.SURVEY_FORM_CHANGE:
            old.append((r["date"].strftime("%m/%d/%Y"), r["licence"], h["name"], sub.strftime("%m/%d/%Y %H:%M"),
                        r["staffed"], r["patients"], r["vent"], r["picu_staffed"] or "", r["picu_patients"]
                        if r["picu_staffed"] else "", r["ms_staffed"], r["ms_patients"]))
        else:
            sid += 1 + int(ids.poisson(2.4))
            aware = sub.replace(tzinfo=FD.TZ[h["tz"]])
            new.append((r["date"].isoformat(), r["licence"], h["name"], "SV%07d" % sid, aware.isoformat(timespec="minutes"),
                        r["staffed"], r["patients"], r["vent"], r["boarders"], r["picu_staffed"] if r["picu_staffed"]
                        else "", r["picu_patients"] if r["picu_staffed"] else "", r["ms_staffed"], r["ms_patients"]))
    n_old = _csv(path_old, ["RptDate", "FacLicense", "FacilityName", "SubmitTime", "AdultICU_Staffed",
                            "AdultICU_Census", "AdultICU_Vent", "PICU_Staffed", "PICU_Census", "MedSurg_Staffed",
                            "MedSurg_Census"], old)
    n_new = _csv(path_new, ["report_date", "licence_no", "facility_name", "submission_id", "submitted_at",
                            "staffed_adult_icu_beds", "adult_icu_patients", "adult_icu_ventilated",
                            "adult_icu_held_in_ed", "staffed_peds_icu_beds", "peds_icu_patients",
                            "staffed_med_surg_beds", "med_surg_patients"], new)
    return n_old, n_new


# -------------------------------------------------------------------------------------------------- board
def write_board(B, path):
    ids = np.random.default_rng(W.SEED + 35)          # the board numbers the events of every facility it serves
    eid = 4400000
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        for i, (u, code, unit, status, reason, by, _lic, _kind) in enumerate(B):
            eid += 1 + int(ids.poisson(1.6))
            fh.write(json.dumps({"event_id": "E%08d" % eid, "logged_utc":
                                 u.strftime("%Y-%m-%dT%H:%M:%SZ"), "facility_code": code, "unit": unit,
                                 "status": status, "reason": reason, "set_by": by}, separators=(",", ":")) + "\n")
    return len(B)


# ------------------------------------------------------------------------------------------ transfer center
def carried(hosp, R):
    """the old system's open cases at the 5 January cut-over that the new system took in: accepted, not
    cancelled, the acceptance in the 18 hours before the cut-over; then, where a region that sent a patient out of
    region in the window has no such case, its latest out-of-region acceptance before the cut-over. A case a
    hospital declined for want of a bed before another accepted it had been settled at the call desk and was not
    carried."""
    by = {h["licence"]: h for h in hosp}
    out = []
    nobed = lambda r: any(c["answer"] == "Declined" and c["why"] in ("bed", "staff") for c in r["calls"])
    for r in R:
        if r["opened"] >= W.TC_CUTOVER or not r["placed"] or r["cancelled"] or nobed(r):
            continue
        acc = [c for c in r["calls"] if c["answer"] == "Accepted"][-1]
        if W.TC_CUTOVER - dt.timedelta(hours=18) <= acc["t"] < W.TC_CUTOVER:
            out.append(r)
    have = {r["sender"]["region"] for r in out if r["kind"] == "Adult ICU"
            and by[[c for c in r["calls"] if c["answer"] == "Accepted"][-1]["to"]]["region"] != r["sender"]["region"]}
    for code in W.RCODE:
        if code in have:
            continue
        cands = [r for r in R if r["kind"] == "Adult ICU" and r["placed"] and not r["cancelled"] and not nobed(r)
                 and r["opened"] < W.TC_CUTOVER and r["sender"]["region"] == code and r not in out
                 and by[[c for c in r["calls"] if c["answer"] == "Accepted"][-1]["to"]]["region"] != code]
        if cands:
            out.append(max(cands, key=lambda r: r["opened"]))
    return sorted(out, key=lambda r: r["opened"])


def _typed(r, alias_by_fid, rng):
    names = alias_by_fid[r["sender"]["fid"]]
    return names[int(rng.integers(0, len(names)))]


def write_transfers(hosp, R, path_old, path_new):
    from openpyxl import Workbook
    by = {h["licence"]: h for h in hosp}
    alias_by_fid = {}
    for nm, fid, _c in RO.aliases(hosp):
        alias_by_fid.setdefault(fid, []).append(nm)
    for f in RO.nonmembers():
        alias_by_fid.setdefault(f["fid"], []).append(f["name"])
        alias_by_fid[f["fid"]].append(f["name"].upper())
    rng = np.random.default_rng(W.SEED + 52)
    typed = {r["no"]: _typed(r, alias_by_fid, rng) for r in R}
    wb = Workbook()
    wb.remove(wb.active)
    sheets = {}
    n_old = 0
    for r in R:
        if r["opened"] >= W.TC_CUTOVER:
            continue
        key = r["opened"].strftime("%b %Y")
        if key not in sheets:
            ws = wb.create_sheet(key)
            ws.append(["Request #", "Opened (CT)", "Sending facility", "Patient ref", "Level of care", "Call",
                       "Facility called", "Called at (CT)", "Response", "Reason"])
            sheets[key] = ws
        for i, c in enumerate(r["calls"], 1):
            code = ""
            if c["answer"] == "Declined":
                code = (FD.REASONS_OLD if c["t"] < dt.datetime(2025, 12, 1) else FD.REASONS_NEW)[c["why"]]
            sheets[key].append([r["no"], r["opened"], typed[r["no"]], r["patient"], CENTER_LEVEL_OLD[r["kind"]], i,
                                by[c["to"]]["name"] if c["to"] else None, c["t"], c["answer"], code or None])
            n_old += 1
    for ws in wb.worksheets:
        ws.column_dimensions["A"].width = 11
        ws.column_dimensions["B"].width = 17
        ws.column_dimensions["C"].width = 34
        ws.column_dimensions["G"].width = 34
        ws.column_dimensions["H"].width = 17
        for row in ws.iter_rows(min_row=2):
            row[1].number_format = "mm/dd/yyyy hh:mm"
            row[7].number_format = "mm/dd/yyyy hh:mm"
    wb.save(path_old)
    # the new system: requests opened from the cut-over, and the open cases it took in
    rows = []
    seq_id = 1
    ids = {}
    for r in carried(hosp, R):
        acc = [c for c in r["calls"] if c["answer"] == "Accepted"][-1]
        ids[r["no"]] = "R26-%05d" % seq_id
        seq_id += 1
        rows.append((ids[r["no"]], r["no"], _utc(r["opened"]), typed[r["no"]], r["patient"],
                     CENTER_LEVEL_NEW[r["kind"]], 1, by[acc["to"]]["center_code"], _utc(acc["t"]), "ACCEPTED", ""))
    for r in R:
        if r["opened"] < W.TC_CUTOVER:
            continue
        ids[r["no"]] = "R26-%05d" % seq_id
        seq_id += 1
        for i, c in enumerate(r["calls"], 1):
            out = {"Accepted": "ACCEPTED", "Declined": "DECLINED", "Cancelled": "CANCELLED",
                   "No placement": "UNPLACED"}[c["answer"]]
            rows.append((ids[r["no"]], "", _utc(r["opened"]), typed[r["no"]], r["patient"], CENTER_LEVEL_NEW[r["kind"]],
                         i, by[c["to"]]["center_code"] if c["to"] else "", _utc(c["t"]), out,
                         FD.REASONS_NEW[c["why"]] if c["answer"] == "Declined" else ""))
    n_new = _csv(path_new, ["request_id", "legacy_case", "opened_utc", "sending_facility_text", "patient_ref",
                            "care_level", "seq", "called_facility_id", "called_utc", "outcome", "decline_code"], rows)
    return n_old, n_new, typed, ids


def _utc(ct_naive):
    return ct_naive.replace(tzinfo=FD.CT).astimezone(FD.UTC).strftime("%Y-%m-%dT%H:%M:%SZ")


# ------------------------------------------------------------------------------------------ the distractors
def write_federal(hosp, S, path):
    """the federal hospital respiratory reporting extract: each hospital's seven-day averages for the weeks ending
    Saturday, from its own midnight census (not the alliance's 07:00 counts)"""
    rng = np.random.default_rng(W.SEED + 53)
    latest = {}
    for r in S:
        k = (r["licence"], r["date"])
        if k not in latest or r["submitted"] > latest[k]["submitted"]:
            latest[k] = r
    rows = []
    week_end = dt.date(2025, 12, 6)
    while week_end <= dt.date(2026, 1, 17):
        for h in sorted(hosp, key=lambda h: h["name"]):
            days = [week_end - dt.timedelta(days=i) for i in range(7)]
            vals = [latest[(h["licence"], d)] for d in days if (h["licence"], d) in latest]
            if not vals:
                continue
            st = np.mean([v["staffed"] for v in vals])
            oc = np.mean([max(0, v["patients"] - rng.integers(0, 2)) for v in vals])
            rows.append((week_end.isoformat(), h["name"], h["county"], "TN", "%.1f" % st, "%.1f" % oc,
                         "%.1f" % np.mean([v["vent"] for v in vals]), len(vals)))
        week_end += dt.timedelta(days=7)
    return _csv(path, ["week_ending", "hospital_name", "county", "state", "staffed_adult_icu_beds_7_day_avg",
                       "staffed_adult_icu_occupied_7_day_avg", "adult_icu_ventilated_7_day_avg", "days_reported"], rows)


def write_history(path):
    """the 2024-25 season's reviews: the committee met on Tuesdays from 3 December to 25 March, not in the two
    holiday weeks; a region under an activation still running past the next Monday is not decided again"""
    rng = np.random.default_rng(W.SEED + 54)
    rows = []
    until = {code: None for code in W.RCODE}
    d = dt.date(2024, 12, 3)
    while d <= dt.date(2025, 3, 25):
        if d in (dt.date(2024, 12, 24), dt.date(2024, 12, 31)):
            d += dt.timedelta(days=7)
            continue
        monday = d + dt.timedelta(days=6)
        for code, name, *_ in W.REGIONS:
            if until[code] is not None and until[code] > monday:
                rows.append((d.isoformat(), code, name, "in force", "", ""))
                continue
            peak = np.exp(-(((d - dt.date(2025, 1, 21)).days) / 28.0) ** 2)
            act = rng.random() < 0.06 + 0.5 * peak * (1.3 if code in ("UC", "WT", "SC") else 0.8)
            if act:
                until[code] = monday + dt.timedelta(days=27)
                rows.append((d.isoformat(), code, name, "activated", monday.isoformat(), until[code].isoformat()))
            else:
                rows.append((d.isoformat(), code, name, "not activated", "", ""))
        d += dt.timedelta(days=7)
    return _csv(path, ["review_date", "region_code", "region_name", "decision", "activation_from", "activation_to"], rows)


# ------------------------------------------------------------------------------------- dictionary, manifest
DICTIONARY = [
    ("compact_regions_counties.csv", "county", "Tennessee county", "name"),
    ("compact_regions_counties.csv", "county_fips", "federal county code", "5 digits"),
    ("compact_regions_counties.csv", "region_code", "the county's surge region (the compact's Annex A)", "NE, ET, SE, UC, MS, SC, WT, MD"),
    ("compact_regions_counties.csv", "region_name", "the region's name", "text"),
    ("facility_crosswalk.csv", "licence_no", "the hospital's state licence number; the census survey's key", "5 digits"),
    ("facility_crosswalk.csv", "facility_name", "the hospital's name; the old transfer call system's name for it", "text"),
    ("facility_crosswalk.csv", "county", "the county the hospital stands in", "name"),
    ("facility_crosswalk.csv", "county_fips", "federal county code", "5 digits"),
    ("facility_crosswalk.csv", "time_zone", "the hospital's time zone", "IANA name"),
    ("facility_crosswalk.csv", "board_code", "the bed-status board's facility code", "F and 4 digits"),
    ("facility_crosswalk.csv", "board_code_previous", "the facility's earlier code on the bed-status board", "F and 4 digits"),
    ("facility_crosswalk.csv", "center_code", "the transfer center's directory code for the hospital", "TC and 3 digits"),
    ("facility_crosswalk.csv", "adult_icu_beds_licensed", "adult ICU beds on the hospital's licence (not all staffed)", "beds"),
    ("facility_crosswalk.csv", "member_since", "the year the hospital joined the alliance", "year"),
    ("transfer_center_name_aliases.csv", "typed_name", "a facility name as the center's staff type it for a sending facility", "text"),
    ("transfer_center_name_aliases.csv", "facility_id", "the licence number of a compact hospital, or the center's id for another facility", "text"),
    ("transfer_center_name_aliases.csv", "facility_name", "the facility's name", "text"),
    ("transfer_center_name_aliases.csv", "county", "the facility's county (Tennessee facilities)", "name"),
    ("transfer_center_name_aliases.csv", "state", "the facility's state", "2 letters"),
    ("transfer_center_name_aliases.csv", "time_zone", "the facility's time zone", "IANA name"),
    ("transfer_center_name_aliases.csv", "compact_member", "Y for a compact hospital", "Y or N"),
    ("transfer_center_reason_codes.csv", "code", "a reason code a declining hospital gives", "text"),
    ("transfer_center_reason_codes.csv", "meaning", "what the code means", "text"),
    ("transfer_center_reason_codes.csv", "in_use_from", "first day the code was in use", "date"),
    ("transfer_center_reason_codes.csv", "in_use_to", "last day the code was in use (blank: in use)", "date"),
    ("capacity_survey_2025-11.csv", "RptDate", "the date the report is for", "MM/DD/YYYY"),
    ("capacity_survey_2025-11.csv", "FacLicense", "the hospital's state licence number", "5 digits"),
    ("capacity_survey_2025-11.csv", "FacilityName", "the hospital's name", "text"),
    ("capacity_survey_2025-11.csv", "SubmitTime", "when the hospital submitted, its own time", "MM/DD/YYYY HH:MM"),
    ("capacity_survey_2025-11.csv", "AdultICU_Staffed", "staffed adult ICU beds at 07:00", "beds"),
    ("capacity_survey_2025-11.csv", "AdultICU_Census", "adult ICU patients at 07:00", "patients"),
    ("capacity_survey_2025-11.csv", "AdultICU_Vent", "adult ICU patients on ventilators at 07:00", "patients"),
    ("capacity_survey_2025-11.csv", "PICU_Staffed", "staffed pediatric ICU beds at 07:00 (blank: no pediatric ICU)", "beds"),
    ("capacity_survey_2025-11.csv", "PICU_Census", "pediatric ICU patients at 07:00", "patients"),
    ("capacity_survey_2025-11.csv", "MedSurg_Staffed", "staffed medical-surgical beds at 07:00", "beds"),
    ("capacity_survey_2025-11.csv", "MedSurg_Census", "medical-surgical patients at 07:00", "patients"),
    ("capacity_survey_2025-12-01_2026-01-19.csv", "report_date", "the date the report is for", "YYYY-MM-DD"),
    ("capacity_survey_2025-12-01_2026-01-19.csv", "licence_no", "the hospital's state licence number", "5 digits"),
    ("capacity_survey_2025-12-01_2026-01-19.csv", "facility_name", "the hospital's name", "text"),
    ("capacity_survey_2025-12-01_2026-01-19.csv", "submission_id", "the survey's id for the submission", "SV and 7 digits"),
    ("capacity_survey_2025-12-01_2026-01-19.csv", "submitted_at", "when the hospital submitted, its own time with its UTC offset", "ISO 8601"),
    ("capacity_survey_2025-12-01_2026-01-19.csv", "staffed_adult_icu_beds", "staffed adult ICU beds at 07:00", "beds"),
    ("capacity_survey_2025-12-01_2026-01-19.csv", "adult_icu_patients", "adult ICU patients at 07:00", "patients"),
    ("capacity_survey_2025-12-01_2026-01-19.csv", "adult_icu_ventilated", "adult ICU patients on ventilators at 07:00", "patients"),
    ("capacity_survey_2025-12-01_2026-01-19.csv", "adult_icu_held_in_ed", "patients admitted to the adult ICU and held in the emergency department at 07:00", "patients"),
    ("capacity_survey_2025-12-01_2026-01-19.csv", "staffed_peds_icu_beds", "staffed pediatric ICU beds at 07:00 (blank: no pediatric ICU)", "beds"),
    ("capacity_survey_2025-12-01_2026-01-19.csv", "peds_icu_patients", "pediatric ICU patients at 07:00", "patients"),
    ("capacity_survey_2025-12-01_2026-01-19.csv", "staffed_med_surg_beds", "staffed medical-surgical beds at 07:00", "beds"),
    ("capacity_survey_2025-12-01_2026-01-19.csv", "med_surg_patients", "medical-surgical patients at 07:00", "patients"),
    ("bed_status_events_2025-11-01_2026-01-20.jsonl", "event_id", "the board's id for the event", "E and 8 digits"),
    ("bed_status_events_2025-11-01_2026-01-20.jsonl", "logged_utc", "when the status was set or confirmed", "ISO 8601, UTC"),
    ("bed_status_events_2025-11-01_2026-01-20.jsonl", "facility_code", "the board's facility code", "F and 4 digits"),
    ("bed_status_events_2025-11-01_2026-01-20.jsonl", "unit", "the unit as the hospital named it on the board; a hospital's adult ICU appears as Adult ICU, ICU, MICU/SICU, Critical Care Unit, Adult Critical Care or Intensive Care", "text"),
    ("bed_status_events_2025-11-01_2026-01-20.jsonl", "status", "the status the hospital set for the unit: Open, Limited (takes patients with restrictions) or Closed (takes no transfer); an emergency department's status is Open, Advisory or Divert", "text"),
    ("bed_status_events_2025-11-01_2026-01-20.jsonl", "reason", "the hospital's note on the status", "free text"),
    ("bed_status_events_2025-11-01_2026-01-20.jsonl", "set_by", "who set or confirmed the status", "role"),
    ("transfer_center_calls_2025-11-01_2026-01-04.xlsx", "Request #", "the old call system's request number", "integer"),
    ("transfer_center_calls_2025-11-01_2026-01-04.xlsx", "Opened (CT)", "when the request was opened, Central time", "date and time"),
    ("transfer_center_calls_2025-11-01_2026-01-04.xlsx", "Sending facility", "the facility asking for the placement, as typed", "text"),
    ("transfer_center_calls_2025-11-01_2026-01-04.xlsx", "Patient ref", "the center's reference for the patient", "P and 7 digits"),
    ("transfer_center_calls_2025-11-01_2026-01-04.xlsx", "Level of care", "the level of care asked for", "Adult ICU, Step-down/PCU, Pediatric ICU"),
    ("transfer_center_calls_2025-11-01_2026-01-04.xlsx", "Call", "the call's order within the request", "integer"),
    ("transfer_center_calls_2025-11-01_2026-01-04.xlsx", "Facility called", "the hospital called", "facility name"),
    ("transfer_center_calls_2025-11-01_2026-01-04.xlsx", "Called at (CT)", "when the answer was logged, Central time", "date and time"),
    ("transfer_center_calls_2025-11-01_2026-01-04.xlsx", "Response", "Accepted, Declined, Cancelled (an acceptance the sending facility cancelled) or No placement", "text"),
    ("transfer_center_calls_2025-11-01_2026-01-04.xlsx", "Reason", "the declining hospital's reason code", "code"),
    ("transfer_center_calls_2026-01-05_2026-01-19.csv", "request_id", "the new call system's request id", "R26- and 5 digits"),
    ("transfer_center_calls_2026-01-05_2026-01-19.csv", "legacy_case", "the previous call system's request number, where the case has one", "integer"),
    ("transfer_center_calls_2026-01-05_2026-01-19.csv", "opened_utc", "when the request was opened", "ISO 8601, UTC"),
    ("transfer_center_calls_2026-01-05_2026-01-19.csv", "sending_facility_text", "the facility asking for the placement, as typed", "text"),
    ("transfer_center_calls_2026-01-05_2026-01-19.csv", "patient_ref", "the center's reference for the patient", "P and 7 digits"),
    ("transfer_center_calls_2026-01-05_2026-01-19.csv", "care_level", "the level of care asked for", "ADULT_ICU, STEP_DOWN, PEDS_ICU"),
    ("transfer_center_calls_2026-01-05_2026-01-19.csv", "seq", "the call's order within the request", "integer"),
    ("transfer_center_calls_2026-01-05_2026-01-19.csv", "called_facility_id", "the hospital called, by the center's directory code", "TC and 3 digits"),
    ("transfer_center_calls_2026-01-05_2026-01-19.csv", "called_utc", "when the answer was logged", "ISO 8601, UTC"),
    ("transfer_center_calls_2026-01-05_2026-01-19.csv", "outcome", "ACCEPTED, DECLINED, CANCELLED (an acceptance the sending facility cancelled) or UNPLACED", "text"),
    ("transfer_center_calls_2026-01-05_2026-01-19.csv", "decline_code", "the declining hospital's reason code", "code"),
    ("federal_hospital_respiratory_weekly_2025-12_2026-01.csv", "week_ending", "the Saturday the week ends on", "YYYY-MM-DD"),
    ("federal_hospital_respiratory_weekly_2025-12_2026-01.csv", "hospital_name", "the hospital's name as reported federally", "text"),
    ("federal_hospital_respiratory_weekly_2025-12_2026-01.csv", "county", "the hospital's county", "name"),
    ("federal_hospital_respiratory_weekly_2025-12_2026-01.csv", "state", "the hospital's state", "2 letters"),
    ("federal_hospital_respiratory_weekly_2025-12_2026-01.csv", "staffed_adult_icu_beds_7_day_avg", "the week's average of staffed adult ICU beds at the hospital's midnight census", "beds"),
    ("federal_hospital_respiratory_weekly_2025-12_2026-01.csv", "staffed_adult_icu_occupied_7_day_avg", "the week's average of occupied staffed adult ICU beds at the midnight census", "beds"),
    ("federal_hospital_respiratory_weekly_2025-12_2026-01.csv", "adult_icu_ventilated_7_day_avg", "the week's average of ventilated adult ICU patients", "patients"),
    ("federal_hospital_respiratory_weekly_2025-12_2026-01.csv", "days_reported", "days of the week the hospital reported", "days"),
    ("export_manifest_2026-01-20.csv", "file", "the file", "file name"),
    ("export_manifest_2026-01-20.csv", "holds", "what the file holds", "text"),
    ("export_manifest_2026-01-20.csv", "covers", "the period the file covers, where it has one", "text"),
    ("export_manifest_2026-01-20.csv", "rows", "the file's data rows (blank for a document)", "rows"),
    ("export_manifest_2026-01-20.csv", "source", "the system or team the file comes from", "text"),
    ("compact_activations_2024-25.csv", "review_date", "the 2024-25 season's review", "YYYY-MM-DD"),
    ("compact_activations_2024-25.csv", "region_code", "the surge region", "NE, ET, SE, UC, MS, SC, WT, MD"),
    ("compact_activations_2024-25.csv", "region_name", "the region's name", "text"),
    ("compact_activations_2024-25.csv", "decision", "the review's decision for the region: activated, not activated, or in force (an earlier activation still running)", "text"),
    ("compact_activations_2024-25.csv", "activation_from", "the activation's first day", "YYYY-MM-DD"),
    ("compact_activations_2024-25.csv", "activation_to", "the activation's last day", "YYYY-MM-DD"),
]


def write_dictionary(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Fields"
    ws.append(["file", "field", "meaning", "format"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in DICTIONARY:
        ws.append(list(r))
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 90
    ws.column_dimensions["D"].width = 26
    ws.freeze_panes = "A2"
    wb.save(path)
    return len(DICTIONARY)


def write_manifest(path, counts):
    rows = [
        (FILES["compact"], "the Ridgeline ICU Surge Compact, effective 1 October 2025", "", "", "alliance board"),
        (FILES["annex"], "the compact's data annex, version of 3 November 2025", "", "", "alliance data team"),
        (FILES["regions"], "every Tennessee county with its surge region (the compact's Annex A)", "", counts["regions"],
         "alliance data team"),
        (FILES["crosswalk"], "the compact hospitals with their keys in each feed", "", counts["crosswalk"],
         "alliance data team"),
        (FILES["aliases"], "the transfer center's typed names for sending facilities", "", counts["aliases"],
         "transfer center"),
        (FILES["codes"], "the transfer center's decline reason codes", "", counts["codes"], "transfer center"),
        (FILES["survey_old"], "census survey submissions, November form", "2025-11-01 to 2025-11-30",
         counts["survey_old"], "census survey"),
        (FILES["survey_new"], "census survey submissions, December form", "2025-12-01 to 2026-01-19",
         counts["survey_new"], "census survey"),
        (FILES["board"], "every status change and confirmation on the bed-status board", "to 2026-01-20 06:00 Central",
         counts["board"], "bed-status board"),
        (FILES["tc_old"], "the transfer center's calls, old call system (requests opened to 4 January)",
         "2025-11-01 to 2026-01-04", counts["tc_old"], "transfer center"),
        (FILES["tc_new"], "the transfer center's calls, new call system (live 5 January)", "2026-01-05 to 2026-01-19",
         counts["tc_new"], "transfer center"),
        (FILES["dict"], "the fields of the data files", "", counts["dict"], "alliance data team"),
        (FILES["federal"], "the federal hospital respiratory reporting extract, weekly, compact hospitals",
         "weeks ending 2025-12-06 to 2026-01-17", counts["federal"], "federal reporting system"),
        (FILES["history"], "the 2024-25 season's reviews and activations", "2024-12-03 to 2025-03-25",
         counts["history"], "alliance data team")]
    return _csv(path, ["file", "holds", "covers", "rows", "source"], rows)
