"""task456's assertions, run on every build: the input gates, the golden from the bundle against the generator's
truth, every wrong route by name, the model's variants and starts, the margins, every semantic-fork row, the
three-rules block on the clean truth, mess neutrality, the asks' stops and cracker, the separation of the asks from
the main call, the distractors' deletion test and the realism sweep."""
import collections
import csv
import datetime as dt
import json
import os
import re
import shutil
import tempfile

import numpy as np

import conform as CF
import golden as GD
import lca
import methods as M
import pack as P
import world as W

RESULTS = []
G = W.RCODE
UTC = dt.timezone.utc


def ok(label, cond, detail=""):
    RESULTS.append((label, bool(cond), detail))


def tenths(sh):
    return {g: M.tenth(sh[g]) for g in G}


def fit_rows(rows):
    Y, reg, _ = M.arrays(rows)
    return lca.fit(Y, reg, G, True)["share"]


def pdf_text(path):
    from pypdf import PdfReader
    return "\n".join(p.extract_text() or "" for p in PdfReader(path).pages)


def headers(target):
    """every shipped column name, by file"""
    from openpyxl import load_workbook
    out = {}
    for f in sorted(os.listdir(target)):
        p = os.path.join(target, f)
        if f.endswith(".csv"):
            with open(p, encoding="utf-8") as fh:
                out[f] = next(csv.reader(fh))
        elif f.endswith(".jsonl"):
            with open(p, encoding="utf-8") as fh:
                out[f] = list(json.loads(fh.readline()).keys())
        elif f.endswith(".xlsx"):
            wb = load_workbook(p, read_only=True)
            out[f] = sorted({str(c) for ws in wb.worksheets for c in next(ws.iter_rows(max_row=1, values_only=True))})
    return out


def run(target, golden_dir, world=None, meta_path=None):
    RESULTS.clear()
    X = world["golden"] if world and "golden" in world else GD.compute(target)
    pk, gd = X["pk"], X["gd"]
    files = sorted(os.listdir(target))
    meta = json.load(open(meta_path)) if meta_path else {}
    H, Z, F = world["H"], world["Z"], world["F"]
    A0, A1 = W.WINDOW_START, W.WINDOW_END
    gold_set = ["NE", "UC", "SC", "WT"]
    gold_t = {"NE": 50, "ET": 20, "SE": 10, "UC": 50, "MS": 10, "SC": 70, "WT": 80, "MD": 10}

    # ---------------------------------------------------------------- input gates
    ok("gate: ten or more files", len(files) >= 10, "%d files" % len(files))
    fmts = {f.rsplit(".", 1)[1] for f in files}
    ok("gate: three or more formats", len(fmts) >= 3, ", ".join(sorted(fmts)))
    ok("gate: a file of 25,000 rows or more", world["counts"]["board"] >= 25000, world["counts"]["board"])
    dis = meta.get("distractor_files", [])
    ok("gate: two or more distractors in metadata.json", len(dis) >= 2 and all(d in files for d in dis), dis)
    ok("gate: no file name says distractor, answer or trap",
       not any(re.search("distract|answer|known_issue|trap|decoy", f) for f in files))
    ok("gate: every pack file has a source and licence in metadata.json",
       sorted(x["path"] for x in meta.get("files", [])) == files and all(x.get("license") for x in meta["files"]))

    # ---------------------------------------------------------------- the golden, from the bundle, against the truth
    season = pk.hospital_days(W.SEASON_START, W.SEASON_END)
    bad = sum(1 for reg, lic, nm, d, a, b, c in season if (a, b, c) != F[(lic, d)])
    ok("golden: the three feeds conformed from the bundle equal the generator's flags on all 6,000 hospital-days of "
       "the season", bad == 0 and len(season) == 6000, "%d mismatches" % bad)
    ok("golden: 75 compact hospitals, 2,100 hospital-days in the 28 days", len(gd["rows"]) == 2100 and
       sum(X["hospitals"].values()) == 75)
    ok("golden: the compact activates Northeast, Upper Cumberland, South Central and West", gd["regions"] == gold_set,
       gd["regions"])
    ok("golden: shares to the nearest ten per cent 50, 20, 10, 50, 10, 70, 80, 10", X["tenth"] == gold_t, X["tenth"])
    edge = {g: abs((10 * gd["share"][g]) % 1 - 0.5) * 10 for g in G}
    ok("bins: every share sits 1.5 points or more from a ten-per-cent rounding edge", min(edge.values()) >= 1.5,
       {g: round(v, 1) for g, v in edge.items()})
    gap = {g: abs(gd["share"][g] - W.THRESHOLD) * 100 for g in G}
    ok("margins: every region's share sits 8 points or more from one in three (SF5)", min(gap.values()) >= 8,
       {g: round(v, 1) for g, v in gap.items()})
    ok("golden: the fit's feed rates (census 33 and 0.8, board 89 and 35, transfer center 24 and 1.6 per cent of full "
       "and ordinary days)", list(np.round(100 * gd["s"]).astype(int)) == [33, 89, 24] and
       (round(100 * gd["f"][0], 1), round(100 * gd["f"][1]), round(100 * gd["f"][2], 1)) == (0.8, 35, 1.6),
       (gd["s"], gd["f"]))
    truth = {g: float(np.mean([Z[(h["licence"], d)] for h in H if h["region"] == g for d in LTdays(A0, A1)]))
             for g in G}
    ok("golden: the estimate sits within 5 points of every region's true share of full days, and the truth activates "
       "the same four regions", max(abs(gd["share"][g] - truth[g]) for g in G) < 0.05 and M.activated(truth) ==
       gold_set, {g: round(truth[g], 3) for g in G})
    starts = [((0.3, 0.6, 0.2), (0.02, 0.4, 0.02)), ((0.8, 0.95, 0.6), (0.1, 0.1, 0.1)), ((0.6, 0.7, 0.5),
                                                                                             (0.01, 0.3, 0.01))]
    dev = max(max(abs(lca.fit(gd["Y"], gd["reg"], G, True, s0=s0, f0=f0)["share"][g] - gd["share"][g]) for g in G)
              for s0, f0 in starts)
    ok("EM: three starts reach the same fit (closure row 14)", dev < 1e-5, "%.1e" % dev)
    sw = lca.fit(gd["Y"], gd["reg"], G, True, s0=(0.01, 0.3, 0.01), f0=(0.3, 0.9, 0.2))
    ok("EM: the swapped start reaches the same fit with the classes' labels swapped (the full class is the one every "
       "feed flags more often)", abs(sw["loglik"] - gd["loglik"]) < 1e-6 and
       max(abs(1 - sw["share"][g] - gd["share"][g]) for g in G) < 1e-4 and all(gd["s"] > gd["f"]))

    # ---------------------------------------------------------------- the wrong routes, by name
    R = X["rules"]
    every = list(G)
    ok("R0 (default): a day is full when any feed shows it: all eight regions",
       R["any feed"]["regions"] == every and min(R["any feed"]["share"].values()) > 0.40,
       {g: round(v, 3) for g, v in R["any feed"]["share"].items()})
    ok("R1: two of the three feeds agree: no region", R["two of three"]["regions"] == [],
       {g: round(v, 3) for g, v in R["two of three"]["share"].items()})
    ok("R1': the board only with the census or a decline: West alone",
       R["C+BC+A+AC+AB+ABC"]["regions"] == ["WT"], {g: round(v, 3) for g, v in R["C+BC+A+AC+AB+ABC"]["share"].items()})
    ok("R2: the board alone: all eight regions", R["B+BC+AB+ABC"]["regions"] == every)
    ok("R2': the census alone: no region", R["A+AC+AB+ABC"]["regions"] == [])
    ok("R2'': the transfer center's declines alone: no region", R["C+BC+AC+ABC"]["regions"] == [])
    outcomes = collections.Counter(tuple(v["regions"]) for v in R.values())
    ok("ladder: all eighteen monotone day-level rules miss the four (five give all eight, eleven none, two West)",
       len(R) == 18 and outcomes == {tuple(every): 5, (): 11, ("WT",): 2}, dict(outcomes))
    V = M.variants(pk, gd)
    ml = V["most likely state, one statewide share"]
    ok("R3: the model's most likely state per day counted, one statewide share: no region", M.activated(ml) == [],
       {g: round(v, 3) for g, v in ml.items()})
    mlr = V["most likely state, by region"]
    wrong = [g for g in G if M.tenth(mlr[g]) != gold_t[g]]
    ok("R3 by region: the most likely state counted keeps the four regions with four of the eight shares wrong "
       "(Northeast, East, Upper Cumberland, South Central)", M.activated(mlr) == gold_set and
       wrong == ["NE", "ET", "UC", "SC"], wrong)
    okA = ~np.isnan(gd["Y"][:, 0])
    lpr = {}
    for g in G:
        y = gd["Y"][gd["reg"] == g]
        o = ~np.isnan(y[:, 0])
        lpr[g] = (y[o, 0].sum() * y[o, 2].sum() / (y[o, 0] * y[o, 2]).sum()) / o.sum()
    ok("R2c: two-source capture-recapture on the census and the declines, region by region: six regions (East and "
       "Southeast added on small overlaps and the declines' false flags)", M.activated(lpr) == ["NE", "ET", "SE", "UC",
                                                                                              "SC", "WT"],
       {g: round(v, 3) for g, v in lpr.items()})
    sA = (gd["Y"][okA, 0] * gd["Y"][okA, 2]).sum() / gd["Y"][okA, 2].sum()
    scl = {g: float(np.nanmean(gd["Y"][gd["reg"] == g, 0]) / sA) for g in G}
    ok("R4': the census scaled by its statewide detection rate from its overlap with the declines (the decisive "
       "insight, a crude estimator): the four regions with Northeast and Mid-State wrong (70 and 20 against 50 and 10)",
       M.activated(scl) == gold_set and [g for g in G if M.tenth(scl[g]) != gold_t[g]] == ["NE", "MS"],
       {g: round(v, 3) for g, v in scl.items()})
    ok("position: the default path activates twice the golden's regions and no wrong rung names the four",
       len(R["any feed"]["regions"]) == 2 * len(gold_set) and all(v["regions"] != gold_set for v in R.values()))

    # ---------------------------------------------------------------- variants of the model (C3 corridor) and SF6
    for k in ("hospital-level shares", "season-wide feed rates", "Beta(2,2) priors", "missing census read as no flag"):
        ok("C3 %s: the same four regions and the same shares to the nearest ten per cent" % k,
           M.activated(V[k]) == gold_set and tenths(V[k]) == gold_t, {g: round(v, 3) for g, v in V[k].items()})
    pm = V["one statewide share, posterior means"]
    ok("SF6 one statewide share (posterior means): the same four regions, 7.5 points or more from the line, other "
       "shares (ruled out by article 1.1, strain is regional)", M.activated(pm) == gold_set and
       min(abs(pm[g] - W.THRESHOLD) for g in G) >= 0.075 and tenths(pm) != gold_t,
       {g: round(v, 3) for g, v in pm.items()})
    sep = {}
    for g in G:
        idx = gd["reg"] == g
        sep[g] = lca.fit(gd["Y"][idx], gd["reg"][idx], [g], True)["share"][g]
    ok("SF6 a separate fit per region with its own feed rates: the same four regions, other shares (ruled out by the "
       "annex, every hospital reports through the same systems under the same instructions in every region)", M.activated(sep) == gold_set and
       tenths(sep) != gold_t, {g: round(v, 3) for g, v in sep.items()})
    ok("SF6 each region's share is the latent prevalence: every day-level rule gives another call",
       all(v["regions"] != gold_set for v in R.values()))

    # ---------------------------------------------------------------- the semantic fork table, row by row
    feed_heads = headers(target)
    ok("SF1 a hospital's region: no feed carries a region; the crosswalk's county and annex A place every hospital",
       not any(re.search("region", h, re.I) for f in (P.FILES["survey_old"], P.FILES["survey_new"], P.FILES["board"],
                                                     P.FILES["tc_old"], P.FILES["tc_new"]) for h in feed_heads[f])
       and all(pk.region[lic] in G for lic in pk.hosp))
    pk.clock("UTC")
    utc = fit_rows(pk.hospital_days(A0, A1))
    pk.clock("Central")
    cen = fit_rows(pk.hospital_days(A0, A1))
    pk.clock(None)
    ok("SF2 a hospital-day on UTC days or on Central days for every hospital: the same four regions (article 2 fixes "
       "the hospital's own time)", M.activated(utc) == gold_set and M.activated(cen) == gold_set,
       ({g: round(v, 3) for g, v in utc.items()}, {g: round(v, 3) for g, v in cen.items()}))
    ok("SF3 a full day read as the board's Closed: all eight regions (article 2.4 and the annex rule it out)",
       R["B+BC+AB+ABC"]["regions"] == every)
    sh = fit_rows(pk.hospital_days(A0 + dt.timedelta(days=1), A1 + dt.timedelta(days=1)))
    ok("SF4 the 28 days to the review's eve (23 December to 19 January): the same four regions",
       M.activated(sh) == gold_set, {g: round(v, 3) for g, v in sh.items()})
    ok("SF5 more than one in three: the same four regions (no share within 8 points of the line)",
       [g for g in G if gd["share"][g] > W.THRESHOLD] == gold_set)
    called = {c["called"] for c in pk.calls() if c["outcome"] != "UNPLACED"}
    codes_on_board = {json.loads(line)["facility_code"] for line in open(os.path.join(target, P.FILES["board"]),
                                                                          encoding="utf-8")}
    ok("SF7 compact hospitals: every board code and every called hospital is one of the crosswalk's 75; the "
       "hospital-days are theirs alone", None not in called and codes_on_board <= set(pk.by_new) and
       len(pk.hosp) == 75)
    peds = CF.ICU_LABELS | {"PICU", "Pediatric ICU", "NICU", "Neonatal ICU"}
    s1 = fit_rows(pk.hospital_days(A0, A1, labels=peds))
    s2 = fit_rows(pk.hospital_days(A0, A1, levels=("ADULT_ICU", "STEP_DOWN", "PEDS_ICU")))
    ok("SF8 pediatric and neonatal ICU lines on the board, or step-down and pediatric declines: the same four regions "
       "and shares", M.activated(s1) == gold_set and tenths(s1) == gold_t and M.activated(s2) == gold_set and
       tenths(s2) == gold_t)
    s3 = fit_rows(pk.hospital_days(A0, A1, latest=False))
    ok("SF9 a day's first census submission instead of its latest: the same four regions and shares",
       M.activated(s3) == gold_set and tenths(s3) == gold_t, {g: round(v, 3) for g, v in s3.items()})
    s4 = fit_rows(pk.hospital_days(A0, A1, codes=None))
    ok("SF10 any decline, whatever its reason: the same four regions and shares", M.activated(s4) == gold_set and
       tenths(s4) == gold_t, {g: round(v, 3) for g, v in s4.items()})
    held = pk.hospital_days(A0, A1, held=True)
    ok("SF11 the census with the patients held in the emergency department added: every flag the same (patients wait "
       "there only when the unit is full at 07:00)", [r[4:] for r in held] == [r[4:] for r in gd["rows"]])
    legacy_nos = {c_["legacy"] for c_ in pk.calls() if c_["legacy"]}
    kept_new = set()
    for c_ in pk.calls():
        if c_["system"] == "old" and c_["request"] in legacy_nos:
            continue
        if c_["level"] == "ADULT_ICU" and c_["outcome"] == "DECLINED" and c_["code"] in CF.NO_BED and c_["called"]:
            kept_new.add((c_["called"], c_["at"].astimezone(pk.tz[c_["called"]]).date()))
    ok("SF22 a case carried across the cut-over read from the new system's copy instead of the old: every transfer "
       "flag the same (no carried case holds a no-bed decline)", kept_new == pk.declined_days(), len(kept_new))
    s5 = fit_rows(pk.hospital_days(A0, A1, margin=1))
    ok("SF11 the census within one bed of full: the same four regions (other shares; article 2.4 says every staffed "
       "bed)", M.activated(s5) == gold_set, {g: round(v, 3) for g, v in s5.items()})
    cc = fit_rows([r for r in gd["rows"] if r[4] is not None])
    ok("SF12 a day with no census report read as no flag, or dropped: the same four regions and shares",
       tenths(V["missing census read as no flag"]) == gold_t and M.activated(cc) == gold_set and tenths(cc) == gold_t,
       {g: round(v, 3) for g, v in cc.items()})
    ok("SF13 shares rounded once: every variant of the C3 corridor rounds to the same ten per cent (asserted above) "
       "and the golden sits 1.5 points or more from an edge", min(edge.values()) >= 1.5)
    a = X["ask_a"]
    kept = M.ask_a(pk, drop_cancelled=False)
    ok("SF14 placed: an acceptance later cancelled counted as a placement gives another figure",
       kept != a and sum(kept.values()) > sum(a.values()))
    rec = M.ask_a_receiving(pk)
    ok("SF15 the patient's own region read as the accepting hospital's: every region's figure moves",
       all(rec[g] != a[g] for g in G), rec)
    memb = collections.Counter(p["sender_region"] for p in pk.placements(A0, A1) if p["sender_region"] and
                               p["sender_region"] != p["to_region"] and p["sender"] in pk.hosp)
    ok("SF15 member senders only (article 2.2 places every hospital, a member or not, in its county's region): every "
       "region's figure falls", all(memb.get(g, 0) < a[g] for g in G), dict(memb))
    comp_txt = " ".join(pdf_text(os.path.join(target, P.FILES["compact"])).split())
    ok("SF15 article 2.2 files every hospital's region, member or not", "a member or not, belongs to the region" in comp_txt)
    conv = all(M.ask_a(pk, zone=z, on=o) == a for z in (None, "Central") for o in ("opened", "accepted"))
    ok("SF16 in those 28 days: the request's opening or its acceptance, on the sending hospital's clock or Central "
       "time, give the same figures", conv)
    b = X["ask_b"]
    dv = M.ask_b(pk, keep=lambda s, r: s == "Divert")
    ok("SF17 asked ambulances to divert read as the Divert status alone: every region's figure falls",
       all(round(dv[g]) < round(b[g]) for g in G), M.whole(dv))
    pk.clock("UTC")
    bu, cu = M.ask_b(pk), M.ask_c(pk)
    pk.clock(None)
    c = X["ask_c"]
    per_h_b = by_region_rounded(pk, pk.hours(CF.ED_LABELS, M.diversion, A0, A1))
    per_h_c = by_region_rounded(pk, pk.hours(CF.ICU_LABELS, lambda s, r: s == "Closed", W.PRIOR_START, W.PRIOR_END))
    ok("SF18 hours on UTC day edges, or rounded hospital by hospital: other figures",
       M.whole(bu) != M.whole(b) and M.whole(cu) != M.whole(c) and (per_h_b != M.whole(b) or per_h_c != M.whole(c)),
       (M.whole(bu), M.whole(cu)))
    ub = union_hours(pk, CF.ED_LABELS, M.diversion, A0, A1)
    uc = union_hours(pk, CF.ICU_LABELS, lambda s, r: s == "Closed", W.PRIOR_START, W.PRIOR_END)
    ok("SF18 the hours in which any of the region's hospitals was diverting or closed, instead of the hours added "
       "over its hospitals: every closed-hours figure and seven of the eight diversion figures fall (Upper "
       "Cumberland's departments never diverted at once); the prompt adds the hours of each of its hospitals",
       all(round(ub[g]) <= round(b[g]) for g in G) and sum(round(ub[g]) < round(b[g]) for g in G) == 7 and
       all(round(uc[g]) < round(c[g]) for g in G), (M.whole(ub), M.whole(uc)))
    cl = M.ask_c(pk, keep=lambda s, r: s in ("Closed", "Limited"))
    ok("SF19 shown closed read as Closed or Limited: every region's figure rises", all(cl[g] > c[g] + 50 for g in G))
    alt = M.by_region(pk, pk.hours(CF.ICU_LABELS, lambda s, r: s == "Closed", W.REVIEW - dt.timedelta(days=28),
                                   W.REVIEW - dt.timedelta(days=1)))
    ok("SF20 the 28 days before the review instead of before the window: every region's figure moves",
       all(round(alt[g]) != round(c[g]) for g in G))

    # ---------------------------------------------------------------- the three rules, executed on the clean truth
    big = {P.FILES["board"]: world["counts"]["board"], P.FILES["survey_new"]: world["counts"]["survey_new"],
           P.FILES["tc_old"]: world["counts"]["tc_old"], P.FILES["tc_new"]: world["counts"]["tc_new"]}
    ok("R1 the field's raw material ships: every census submission, every board status event and every transfer "
       "call at the systems' own grain", all(v > 1000 for v in big.values()), big)
    allh = [h for hs in feed_heads.values() for h in hs]
    ok("R1 no shipped column holds a full day, a flag, a share, a probability or a rate",
       not any(re.search(r"(?i)full|flag|share|prob|posterior|prevalence|rate\b", h) for h in allh),
       [h for h in allh if re.search(r"(?i)full|flag|share|prob", h)])
    ok("R1 the golden's decisive computation reads the three feeds' flags on all 2,100 hospital-days",
       gd["Y"].shape == (2100, 3) and np.nansum(gd["Y"], 0).min() > 100)
    days28 = LTdays(A0, A1)
    clean = [(h["region"], h["licence"], h["name"], d) + F[(h["licence"], d)] for h in H for d in days28]
    Yc, rc, _ = M.arrays(clean)
    any_rule = [r for r in lca.monotone_rules() if lca.rule_name(r) == "any feed"][0]
    two_rule = [r for r in lca.monotone_rules() if lca.rule_name(r) == "two of three"][0]
    ok("R3a analyst test: on the clean truth, a day full when any feed shows it activates all eight regions",
       M.activated(lca.rule_share(Yc, rc, G, any_rule)) == every)
    mc = lca.fit(Yc, rc, G, True)["share"]
    ok("R3b the decisive step alone: the same clean flags through the latent class model activate the four",
       M.activated(mc) == gold_set)
    lower = {nm: M.activated(lca.rule_share(Yc, rc, G, r)) for r in lca.monotone_rules()
             for nm in [lca.rule_name(r)] if nm in ("two of three", "C+BC+A+AC+AB+ABC", "B+BC+AB+ABC",
                                                     "A+AC+AB+ABC")}
    ok("R3b each lower rung's single correction on the clean truth stays wrong (two of three, the corroborated "
       "board, the board alone, the census alone)", all(v != gold_set for v in lower.values()) and len(lower) == 4,
       lower)
    ok("R3b two of three on the clean truth activates none", M.activated(lca.rule_share(Yc, rc, G, two_rule)) == [])
    # R2: every Layer 1 item mishandled leaves the call (the fork rows above cover clocks, labels, codes, versions)
    lg = fit_rows(pk.hospital_days(A0, A1, logged=True))
    ok("R2 the board's statuses carried in from the day before ignored (a day counts only with a Closed logged on "
       "it): the same four regions and shares", M.activated(lg) == gold_set and tenths(lg) == gold_t)
    strict = fit_rows(pk.hospital_days(A0, A1, labels={"Adult ICU"}))
    ok("R2 the board's unit labels read narrowly (only lines labelled Adult ICU): the same four regions (other shares; "
       "the dictionary lists the six labels)", M.activated(strict) == gold_set,
       {g: round(v, 3) for g, v in strict.items()})
    win_old = sum(1 for r in csv.DictReader(open(os.path.join(target, P.FILES["survey_old"]), encoding="utf-8"))
                  if A0 <= dt.datetime.strptime(r["RptDate"], "%m/%d/%Y").date())
    ok("R2 the survey's November form (renamed columns) holds no report in the 28 days", win_old == 0)
    win_codes = collections.Counter(c["code"] for c in pk.calls() if c["outcome"] == "DECLINED" and c["called"] and
                                    A0 <= c["at"].astimezone(pk.tz[c["called"]]).date() <= A1)
    ok("R2 the decline code change: no NB or NS decline falls in the 28 days (CAP throughout)",
       win_codes["NB"] == 0 and win_codes["NS"] == 0 and win_codes["CAP"] > 100, dict(win_codes))
    ev = pk.events()
    first = {}
    for t, lic, unit, status, reason in ev:
        if unit in CF.ICU_LABELS and t < dt.datetime(2025, 12, 22, 5, 0, tzinfo=UTC):
            first[(lic, unit)] = t
    ok("R2 the board's renumbering: every adult ICU line's status in force at the window's start was logged after the "
       "migration, and no event in the window carries an earlier code", min(first.values()) > pk.migration)
    ok("R2 the transfer center's names: every called hospital in both exports maps to the crosswalk", None not in
       {c["called"] for c in pk.calls() if c["outcome"] in ("ACCEPTED", "DECLINED", "CANCELLED")})
    carried = [c for c in pk.calls() if c["legacy"]]
    ok("R2 the cut-over: the cases the new system took in carry no decline (zero rows in the call's population)",
       len(carried) > 0 and all(c["outcome"] == "ACCEPTED" for c in carried), len(carried))
    ok("R3c no plant: no document states a share, a rule, an estimator or a set of regions",
       not re.search(r"(?i)latent|maximum likelihood|two of the three|any feed|per cent of (?:the )?(?:region|hospital"
                     r")-days|\b\d{1,2}(?:\.\d)?\s?%", pdf_text(os.path.join(target, P.FILES["compact"])) +
                     pdf_text(os.path.join(target, P.FILES["annex"]))))
    sv = [r for r in csv.DictReader(open(os.path.join(target, P.FILES["survey_new"]), encoding="utf-8"))]
    by_h = collections.defaultdict(list)
    for r in sorted(sv, key=lambda r: (r["licence_no"], r["report_date"], r["submitted_at"])):
        by_h[r["licence_no"]].append((r["staffed_adult_icu_beds"], r["adult_icu_patients"], r["adult_icu_ventilated"],
                                      r["med_surg_patients"]))
    blocks = collections.Counter(tuple(v[i:i + 5]) for v in by_h.values() for i in range(len(v) - 4))
    ok("R3c no copied block: no five consecutive census reports repeat anywhere", max(blocks.values()) == 1)

    # ---------------------------------------------------------------- ask A: out-of-region placements
    gold_a = {"NE": 70, "ET": 40, "SE": 32, "UC": 45, "MS": 49, "SC": 65, "WT": 58, "MD": 28}
    ok("A golden: placements outside the patient's region by sending region, 387 in all", a == gold_a and
       sum(a.values()) == 387, a)
    nat = M.ask_a(pk, carried_once=False, drop_cancelled=False)
    crk = M.ask_a(pk, carried_once=False)
    ded = M.ask_a_patient_dedup(pk)
    for nm, v in (("natural read (both exports summed, cancellations kept)", nat),
                  ("cracker (cancellations dropped, carried cases twice)", crk),
                  ("over-corrected (deduplicated on the patient reference)", ded),
                  ("right rule, wrong population (the receiving region)", rec)):
        ok("A stop, %s: every region's figure misses" % nm, all(v[g] != a[g] for g in G), v)
    pl = pk.placements(A0, A1)
    legacy = {c["legacy"] for c in pk.calls() if c["legacy"]}
    co = collections.Counter(p["sender_region"] for p in pl if p["case"] in legacy and p["sender_region"] and
                             p["sender_region"] != p["to_region"])
    ok("A device: every sending region has a carried case placed out of region", all(co[g] >= 1 for g in G), dict(co))
    twice = collections.Counter((p["sender_region"], p["patient"]) for p in pl if p["sender_region"] and
                                p["sender_region"] != p["to_region"])
    ok("SF21 how many times: no patient has two counted placements from one region, so counting placements and "
       "counting patients converge region by region", max(twice.values()) == 1)
    near = 0
    for r in pk.calls():
        if r["seq"] != 1:
            continue
        _, _, tz = pk.sender(r["sender"])
        loc = r["opened"].astimezone(tz)
        for e in (A0, A1 + dt.timedelta(days=1)):
            if abs((loc - dt.datetime.combine(e, dt.time(0), tzinfo=tz)).total_seconds()) < 2 * 3600:
                near += 1
    ok("A: no request opened within two hours of either edge of the 28 days on the sending hospital's clock", near == 0,
       near)
    lv = M.ask_a(pk, level=None)
    ok("A hazard, step-down and pediatric requests counted: every region's figure rises", all(lv[g] > a[g] for g in G))

    # ---------------------------------------------------------------- ask B: emergency department diversion hours
    gold_b = {"NE": 80, "ET": 133, "SE": 132, "UC": 27, "MS": 147, "SC": 46, "WT": 64, "MD": 69}
    ok("B golden: diversion hours by region, 698 in all", M.whole(b) == gold_b and sum(gold_b.values()) == 698,
       M.whole(b))
    allad = M.ask_b(pk, keep=lambda s, r: s in ("Divert", "Advisory"))
    ok("B stop, every Advisory hour counted: every region's figure misses", all(round(allad[g]) > round(b[g]) + 50
                                                                                 for g in G))
    ed = [e for e in ev if e[2] in CF.ED_LABELS]
    sw_t = min(e[0] for e in ed if e[3] == "Advisory" and e[4] == "EMS diversion")
    ok("B device: no Divert after 5 January 06:00 Central, no Advisory with the reason EMS diversion before it, and "
       "the reason on Divert statuses before it", not any(e[3] == "Divert" and e[0] >= sw_t for e in ed) and
       sw_t >= dt.datetime(2026, 1, 5, 12, 0, tzinfo=UTC) and any(e[3] == "Divert" and e[4] == "EMS diversion"
                                                                    for e in ed), sw_t)
    ok("B and C bins: every hours figure sits 0.1 hour or more from a rounding edge, and the totals of the rounded "
       "figures equal the rounded totals", min(abs(v % 1 - 0.5) for v in list(b.values()) + list(c.values())) >= 0.1
       and sum(M.whole(b).values()) == round(sum(b.values())) and sum(M.whole(c).values()) == round(sum(c.values())))

    # ---------------------------------------------------------------- ask C: adult ICU closed hours before the window
    gold_c = {"NE": 752, "ET": 642, "SE": 475, "UC": 357, "MS": 745, "SC": 513, "WT": 616, "MD": 623}
    ok("C golden: adult ICU closed hours, 24 November to 21 December, by region, 4,723 in all",
       M.whole(c) == gold_c and sum(gold_c.values()) == 4723, M.whole(c))
    cur = M.ask_c_current_code(pk)
    ok("C stop, every event joined on the current code: every region's figure misses",
       all(round(cur[g]) != round(c[g]) for g in G), M.whole(cur))
    xw = list(csv.DictReader(open(os.path.join(target, P.FILES["crosswalk"]), encoding="utf-8")))
    ok("C device: the earlier codes are the same code space as the current ones (a join on the current code leaves no "
       "event unmatched) and 70 or more hospitals changed code",
       {r["board_code_previous"] for r in xw} == {r["board_code"] for r in xw} and
       sum(r["board_code_previous"] != r["board_code"] for r in xw) >= 70 and codes_on_board <= set(pk.by_new))

    # ---------------------------------------------------------------- separation of the asks from the main call
    ok("separation: no emergency department line uses an adult ICU label, and no adult ICU line shows Divert or "
       "Advisory", not (CF.ED_LABELS & CF.ICU_LABELS) and not any(e[3] in ("Divert", "Advisory") for e in ev
                                                                  if e[2] in CF.ICU_LABELS))
    ok("separation: the main call reads no ask device row (carried cases accepted only; the window's board events all "
       "on current codes; no diversion status on an adult ICU line)", all(c["outcome"] == "ACCEPTED" for c in carried)
       and min(first.values()) > pk.migration)

    # ---------------------------------------------------------------- distractors: deleted, nothing moves
    tmp = tempfile.mkdtemp()
    try:
        for f in files:
            if f not in dis:
                shutil.copy(os.path.join(target, f), tmp)
        pk2 = CF.Pack(tmp, P.FILES)
        g2 = M.golden(pk2)
        ok("distractors: with both deleted, the golden set, shares and every ask are unchanged",
           g2["regions"] == gold_set and tenths(g2["share"]) == gold_t and M.ask_a(pk2) == a and
           M.whole(M.ask_b(pk2)) == gold_b and M.whole(M.ask_c(pk2)) == gold_c)
    finally:
        shutil.rmtree(tmp)
    hist = list(csv.DictReader(open(os.path.join(target, P.FILES["history"]), encoding="utf-8")))
    ok("distractors: last season's record carries no share, and no file ranks this season's regions",
       not any(re.search("share|full", k) for k in hist[0]))

    # ---------------------------------------------------------------- realism sweep
    card = json.load(open(os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
                                       ".claude", "skills", "fingerprint", "cards", "task456.json")))
    ok("realism: every persona is one of the card's drawn people", set(W.PEOPLE.values()) <= set(card["world"]["people"]))
    ok("realism: no ask's regions share a figure", all(len(set(v.values())) == 8 for v in (a, gold_b, gold_c)))
    ok("realism: no round ask total", all(t % 10 for t in (387, 698, 4723)))
    end = dt.datetime(2026, 1, 20, 12, 0, tzinfo=UTC)
    ok("realism: every board event and every transfer call is logged before the extract (20 January 06:00 Central)",
       max(e[0] for e in ev) < end and max(c["at"] for c in pk.calls()) < end)
    late = [r for r in sv if dt.datetime.fromisoformat(r["submitted_at"]).replace(tzinfo=None) <
            dt.datetime.fromisoformat(r["report_date"] + "T07:00")]
    ok("realism: every census submission comes after the 07:00 count of its date", not late, len(late))
    per_day = collections.Counter(c["opened"].astimezone(CF.CT).date() for c in pk.calls() if c["seq"] == 1)
    ok("realism: daily request counts vary (no constant stream)", len(set(per_day.values())) > 15)
    stamp = {int(os.path.getmtime(os.path.join(target, f))) for f in files}
    ok("realism: every file carries the one export stamp", len(stamp) == 1)
    for f in (P.FILES["compact"], P.FILES["annex"]):
        raw = open(os.path.join(target, f), "rb").read()
        ok("realism: %s carries no generator banner" % f, b"ReportLab" not in raw and b"reportlab" not in raw)
    docs_txt = pdf_text(os.path.join(target, P.FILES["compact"])) + pdf_text(os.path.join(target, P.FILES["annex"]))
    ok("realism: no em or en dash in the documents", "\u2014" not in docs_txt and "\u2013" not in docs_txt)
    comp = pdf_text(os.path.join(target, P.FILES["compact"]))
    annex = pdf_text(os.path.join(target, P.FILES["annex"]))
    ok("one place each: the full day is defined in the compact alone, and the feeds' errors and the same reporting "
       "in every region in the annex alone", "every staffed adult intensive care bed" in comp.replace("\n", " ") and
       "every staffed adult intensive care bed" not in annex.replace("\n", " ") and
       "same instructions" in " ".join(annex.split()) and "same instructions" not in " ".join(comp.split()))

    # ---------------------------------------------------------------- registers (H9) and the design note's figures (H2)
    man = [r["file"] for r in csv.DictReader(open(os.path.join(target, P.FILES["manifest"]), encoding="utf-8"))]
    ok("registers: the manifest names every shipped file but itself, once", sorted(man) == sorted(
        f for f in files if f != P.FILES["manifest"]) and len(man) == len(set(man)))
    data_files = {f for f in files if f.endswith((".csv", ".jsonl", ".xlsx")) and f != P.FILES["dict"]}
    ok("registers: the dictionary describes every data file and every column of each",
       {r[0] for r in P.DICTIONARY} == data_files and all(
           set(feed_heads[f]) == {r[1] for r in P.DICTIONARY if r[0] == f} for f in data_files),
       [f for f in data_files if set(feed_heads[f]) != {r[1] for r in P.DICTIONARY if r[0] == f}])
    two = lambda d: [round(d[g], 2) for g in G]
    ladder = {"any feed": [0.66, 0.46, 0.41, 0.66, 0.45, 0.76, 0.83, 0.44],
              "two of three": [0.24, 0.09, 0.05, 0.24, 0.06, 0.32, 0.32, 0.05],
              "C+BC+A+AC+AB+ABC": [0.28, 0.11, 0.07, 0.26, 0.08, 0.32, 0.36, 0.07],
              "B+BC+AB+ABC": [0.61, 0.44, 0.40, 0.64, 0.42, 0.74, 0.78, 0.42],
              "A+AC+AB+ABC": [0.19, 0.06, 0.04, 0.15, 0.05, 0.21, 0.23, 0.04]}
    ok("design note: the ladder table's shares reproduce at two decimals",
       all(two(R[k]["share"]) == v for k, v in ladder.items()) and
       two(ml) == [0.27, 0.09, 0.06, 0.24, 0.07, 0.32, 0.33, 0.06] and
       two(gd["share"]) == [0.53, 0.18, 0.10, 0.52, 0.12, 0.68, 0.79, 0.11],
       {k: two(R[k]["share"]) for k in ladder})
    ok("design note: the asks' stops reproduce", [nat[g] for g in G] == [77, 44, 33, 47, 53, 71, 68, 29] and
       [crk[g] for g in G] == [74, 41, 33, 46, 51, 67, 60, 29] and [ded[g] for g in G] == [67, 35, 30, 42, 44, 62, 55, 23]
       and [rec[g] for g in G] == [38, 83, 68, 22, 70, 16, 23, 67] and
       [round(dv[g]) for g in G] == [44, 80, 56, 15, 67, 25, 23, 39] and
       [round(allad[g]) for g in G] == [280, 432, 384, 152, 496, 194, 191, 389] and
       sum(round(bu[g]) != round(b[g]) for g in G) == 4 and
       [round(cur[g]) for g in G] == [668, 654, 532, 398, 802, 490, 599, 583] and
       [round(cl[g]) for g in G] == [1138, 1084, 955, 519, 1639, 852, 900, 1421] and
       min(round(cl[g]) - round(c[g]) for g in G) >= 160 and sum(round(cu[g]) != round(c[g]) for g in G) == 7)
    ok("design note: the variants' margins (statewide posterior means 7.8 points from the line; separate fits six "
       "shares wrong)", round(100 * min(abs(pm[g] - W.THRESHOLD) for g in G), 1) == 7.8 and
       sum(M.tenth(sep[g]) != gold_t[g] for g in G) == 6)

    # ---------------------------------------------------------------- the deliverables
    rows = list(csv.DictReader(open(os.path.join(golden_dir, GD.CSV_NAME), encoding="utf-8")))
    same = all((r["census_full_at_0700"], r["board_adult_icu_closed"], r["transfer_center_no_bed_decline"]) ==
               ("" if F[(r["licence_no"], dt.date.fromisoformat(r["date"]))][0] is None else
                str(F[(r["licence_no"], dt.date.fromisoformat(r["date"]))][0]),
                str(F[(r["licence_no"], dt.date.fromisoformat(r["date"]))][1]),
                str(F[(r["licence_no"], dt.date.fromisoformat(r["date"]))][2])) for r in rows)
    ok("deliverable: the conformed file holds the 2,100 hospital-days, each feed's reading equal to the truth's",
       len(rows) == 2100 and same)
    txt = pdf_text(os.path.join(golden_dir, GD.PDF_NAME)).replace("\n", " ")
    ok("deliverable: the notice names the four regions and carries every share and ask figure",
       all(W.RNAME[g] in txt for g in gold_set) and all("%d%%" % v in txt for v in gold_t.values()) and
       all(format(v, ",") in txt for v in list(gold_a.values()) + list(gold_b.values()) + list(gold_c.values())))

    n_ok = sum(1 for _, c_, _ in RESULTS if c_)
    for label, c_, detail in RESULTS:
        if not c_:
            print("FAIL", label, detail)
    print("checks: %d of %d pass" % (n_ok, len(RESULTS)))
    return n_ok == len(RESULTS)


def LTdays(a, b):
    return [a + dt.timedelta(days=i) for i in range((b - a).days + 1)]


def union_hours(pk, labels, keep, start, end):
    """the hours in which at least one of a region's hospitals held a status keep() accepts (each hospital's own day
    edges)"""
    ivs = collections.defaultdict(list)
    for lic, xs in pk.intervals(labels, keep).items():
        lo = dt.datetime.combine(start, dt.time(0, 0), tzinfo=pk.tz[lic])
        hi = dt.datetime.combine(end + dt.timedelta(days=1), dt.time(0, 0), tzinfo=pk.tz[lic])
        for a, b in xs:
            a2, b2 = max(a, lo), min(b, hi)
            if b2 > a2:
                ivs[pk.region[lic]].append((a2, b2))
    out = {}
    for g in G:
        tot, cur = 0.0, None
        for a, b in sorted(ivs[g]):
            if cur is None or a > cur[1]:
                if cur is not None:
                    tot += (cur[1] - cur[0]).total_seconds() / 3600
                cur = [a, b]
            else:
                cur[1] = max(cur[1], b)
        if cur is not None:
            tot += (cur[1] - cur[0]).total_seconds() / 3600
        out[g] = tot
    return out


def by_region_rounded(pk, hours):
    c = collections.defaultdict(int)
    for lic, v in hours.items():
        c[pk.region[lic]] += int(np.floor(v + 0.5))
    return {g: c.get(g, 0) for g in G}
