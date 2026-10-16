"""Read the shipped pack the way the plan defines it: every supermarket's records put into product-months (a GTIN at a
supermarket, a calendar month or a week by the month it ends in, the unit value with GST, returns netted, each resent
week at its latest extract), each line classified to its item; the parallel-year panel; the published rates; and the
ask files (the delivery log, the annual returns, the category map and the household spending estimates)."""
import csv
import datetime as dt
import os
import sqlite3
from collections import defaultdict

import numpy as np

import world as W


def path(target, key):
    return os.path.join(target, W.FILES[key])


def month_key(d):
    return "%04d-%02d" % (d.year, d.month)


MONTH_KEYS = ["%04d-%02d" % ym for ym in W.MONTHS]


def read_classes(target):
    cl = {}
    with open(path(target, "classes"), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            cl.setdefault((r["retailer"], r["gtin"]), []).append((r["valid_from"], r["valid_to"] or "9999-12-31",
                                                                  r["rpi_item"]))
    return cl


def item_of(cl, retailer, gtin, month):
    first = month + "-01"
    for a, b, it in cl.get((retailer, gtin), []):
        if a <= first <= b:
            return it
    return None


def product_months(target):
    """{(retailer, gtin, month): [units, value with GST]} over the parallel year, as the plan defines them."""
    pm = defaultdict(lambda: [0, 0.0])
    # Wrenfold: weekly rows, each week at its latest extract
    latest = {}
    with open(path(target, "wrenfold"), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            k = (r["week_ending"], r["gtin"])
            if k not in latest or r["extract_date"] > latest[k]["extract_date"]:
                latest[k] = r
    for (we, g), r in latest.items():
        m = month_key(dt.date.fromisoformat(we))
        if m in MONTH_KEYS:
            c = pm[(W.RNAME["W"], g, m)]
            c[0] += int(r["units"])
            c[1] += float(r["sales_value"])
    # Gorsedale: fiscal weeks dated through the calendar, EANs back to 13 digits
    cal = {}
    with open(path(target, "calendar"), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            cal[(r["fiscal_year"], int(r["week"]))] = dt.date.fromisoformat(r["week_end"])
    with open(path(target, "gorsedale"), encoding="utf-8") as fh:
        for r in csv.DictReader(fh, delimiter="|"):
            m = month_key(cal[(r["FISCAL_YEAR"], int(r["WEEK"]))])
            if m in MONTH_KEYS:
                c = pm[(W.RNAME["G"], r["EAN"].zfill(13), m)]
                c[0] += int(r["QTY"])
                c[1] += float(r["SALES"])
    # Calderby's: daily store rows through the SKU table, values grossed up for GST
    con = sqlite3.connect(path(target, "calderby"))
    for d, g, u, v in con.execute("SELECT s.sale_date, k.gtin, s.units, s.net_value FROM daily_sales s "
                                  "JOIN sku_gtin k ON k.sku = s.sku"):
        m = d[:7]
        if m in MONTH_KEYS:
            c = pm[(W.RNAME["C"], g, m)]
            c[0] += int(u)
            c[1] += float(v)
    con.close()
    out = {}
    for k, (u, v) in pm.items():
        if k[0] == W.RNAME["C"]:
            v = v * (1 + W.GST)
        out[k] = (u, v)
    return out


def item_arrays(target, pm=None, cl=None, panel=True):
    """For each item, the lines (retailer, gtin) and arrays P, Q, E over the thirteen months: the panel of lines sold in
    every month when panel is True, otherwise every line the item holds."""
    pm = pm if pm is not None else product_months(target)
    cl = cl if cl is not None else read_classes(target)
    lines = defaultdict(dict)
    for (r, g, m), (u, v) in pm.items():
        it = item_of(cl, r, g, m)
        if it in W.INAME:
            lines[it].setdefault((r, g), {})[m] = (u, v)
    out = {}
    for it in W.ICODE:
        keys = sorted(k for k, ms in lines[it].items()
                      if not panel or all(m in ms and ms[m][0] > 0 for m in MONTH_KEYS))
        P = np.full((len(keys), W.NM), np.nan)
        Q = np.zeros((len(keys), W.NM))
        for i, k in enumerate(keys):
            for j, m in enumerate(MONTH_KEYS):
                if m in lines[it][k] and lines[it][k][m][0] > 0:
                    u, v = lines[it][k][m]
                    P[i, j], Q[i, j] = v / u, u
        out[it] = (keys, P, Q, np.nan_to_num(P) * Q)
    return out


def published_rates(target):
    import openpyxl
    wb = openpyxl.load_workbook(path(target, "published"), read_only=True)
    ws = wb["12-month rates"]
    rows = list(ws.iter_rows(values_only=True))
    head = rows[2]
    col = head.index("Sep 2026")
    return {str(r[0]): float(r[col]) for r in rows[3:] if r[0]}


def published_levels(target):
    """Each item's published index levels by month, from the item indices sheet."""
    import openpyxl
    ws = openpyxl.load_workbook(path(target, "published"), read_only=True)["Item indices"]
    rows = list(ws.iter_rows(values_only=True))
    head = rows[3]
    out = {}
    for r in rows[4:]:
        if not r[0]:
            continue
        lv = {}
        for j, h in enumerate(head[2:], start=2):
            d = dt.datetime.strptime(h, "%b %Y")
            lv["%04d-%02d" % (d.year, d.month)] = float(r[j])
        out[str(r[0])] = lv
    return out


def deliveries(target):
    with open(path(target, "deliveries"), encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def kellesay_time(utc_text):
    t = dt.datetime.strptime(utc_text, "%Y-%m-%dT%H:%M:%SZ")
    y = t.year

    def last_sunday(mo):
        d = dt.date(y, mo, 31)
        return d - dt.timedelta(days=(d.weekday() + 1) % 7)
    start = dt.datetime.combine(last_sunday(3), dt.time(1))
    end = dt.datetime.combine(last_sunday(10), dt.time(1))
    return t + dt.timedelta(hours=1 if start <= t < end else 0)


def month_complete(target, local=True, week_rule=True, resends=True):
    """For each supermarket and month of the parallel year, the date the last file carrying the month's records
    arrived: a file carries a month's records when a day of its records falls in the month (calendar days) or, for a
    supermarket's weekly records, when one of its weeks ends in the month."""
    rows = deliveries(target)
    if not resends:
        rows = [r for r in rows if not (r["file_name"].endswith("_R1.csv") or r["file_name"].endswith("_v2.sqlite"))]
    out = defaultdict(dict)
    for r in rows:
        a = dt.date.fromisoformat(r["covers_from"])
        b = dt.date.fromisoformat(r["covers_to"])
        when = (kellesay_time(r["received_utc"]) if local else dt.datetime.strptime(r["received_utc"], "%Y-%m-%dT%H:%M:%SZ")).date()
        if r["retailer"] == W.RNAME["C"] or not week_rule:
            months = set()
            d = a
            while d <= b:
                months.add(month_key(d))
                d += dt.timedelta(days=1)
        else:
            months = set()
            d = a + dt.timedelta(days=6)
            while d <= b:
                months.add(month_key(d))
                d += dt.timedelta(days=7)
        for m in months:
            if m not in MONTH_KEYS:
                continue
            prev = out[r["retailer"]].get(m)
            if prev is None or when > prev:
                out[r["retailer"]][m] = when
    return out


def delivery_figures(target, **kw):
    comp = month_complete(target, **kw)
    res = {}
    for r in (W.RNAME[x] for x in W.RETAILERS):
        lags = []
        for m in MONTH_KEYS:
            y, mo = int(m[:4]), int(m[5:])
            end = dt.date(y + (mo == 12), mo % 12 + 1, 1) - dt.timedelta(days=1)
            lags.append((comp[r][m] - end).days)
        res[r] = {"lags": lags, "average": float(np.mean(lags)), "late": int(sum(1 for x in lags if x > W.DEADLINE_DAYS))}
    return res


def coverage_figures(target, by_name=True, gross_up=True, retail_only=True):
    """Each item's 2025 retail sales at the three supermarkets, with GST, as a share of households' 2025 spending."""
    import openpyxl
    wb = openpyxl.load_workbook(path(target, "returns"), read_only=True)
    cmap = {}
    with open(path(target, "catmap"), encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            cmap[(r["retailer"], r["category_name"] if by_name else r["category_code"])] = r["rpi_item"]
    sales = defaultdict(float)
    for sheet, retailer in ((W.RNAME["W"], W.RNAME["W"]), (W.RNAME["G"], W.RNAME["G"]), (W.RNAME["C"], W.RNAME["C"])):
        rows = list(wb[sheet].iter_rows(values_only=True))[4:]
        for row in rows:
            if row[0] is None:
                continue
            if retailer == W.RNAME["G"]:
                channel, code, name, val = row
                if retail_only and channel != "Retail":
                    continue
            else:
                code, name, val = row
            it = cmap.get((retailer, name if by_name else str(code)))
            if it is None:
                continue
            v = float(val)
            if retailer == W.RNAME["C"] and gross_up:
                v *= 1 + W.GST
            sales[it] += v
    sw = openpyxl.load_workbook(path(target, "spending"), read_only=True)
    spend = {str(r[0]): float(r[2]) for r in list(sw.active.iter_rows(values_only=True))[4:] if r[0]}
    return {it: 100.0 * sales[it] / 1000.0 / spend[it] for it in W.ICODE}
