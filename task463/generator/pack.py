"""Write the pack's data files: the three feeds, Gorsedale's fiscal calendar, Calderby's database, the office's product
classification, the published item indices, the delivery log, the field quotes, the coverage files and the retail
sales index."""
import csv
import datetime as dt
import os
import sqlite3

import numpy as np

import coverage as CV
import feeds as F
import published as PB
import world as W


def write_csv(path, header, rows, delimiter=",", lineterminator="\r\n"):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh, delimiter=delimiter, lineterminator=lineterminator)
        wr.writerow(header)
        for r in rows:
            wr.writerow(r)
    return len(rows)


WREN_COLS = ["week_ending", "gtin", "description", "units", "sales_value", "extract_date"]
GORS_COLS = ["FISCAL_YEAR", "PERIOD", "WEEK", "ITEM_NO", "EAN", "DESCRIPTION", "QTY", "SALES"]
CAL_COLS = ["fiscal_year", "period", "week", "week_start", "week_end"]
CLASS_COLS = ["retailer", "gtin", "description", "rpi_item", "valid_from", "valid_to"]
DELIV_COLS = ["retailer", "file_name", "covers_from", "covers_to", "rows", "received_utc"]
QUOTE_COLS = ["month", "item_code", "outlet_code", "outlet_name", "product_code", "description", "price", "status",
              "collected_on"]
MAP_COLS = ["retailer", "category_code", "category_name", "rpi_item"]


# the office classified the lines on sale before the parallel year in batches by department, summer 2025
CLASSIFIED = {"11": "2025-07-14", "12": "2025-07-28", "56": "2025-08-11", "62": "2025-08-11"}


def classification_rows(ls, seed):
    out = []
    for L in ls:
        name = L["desc"]
        frm = CLASSIFIED[L["item"][:2]] if L["start"] == 0 else "%04d-%02d-01" % W.MONTHS[L["start"]]
        out.append([W.RNAME[L["retailer"]], L["gtin"], name, L["item"], frm, ""])
    # two lines the office moved between items outside the sixteen during 2026
    moved = [r for r in out if r[3] == "1191"][:2]
    for r in moved:
        r[5] = "2026-03-31"
        out.append([r[0], r[1], r[2], "1109", "2026-04-01", ""])
    out.sort(key=lambda r: (r[0], r[3], r[1], r[4]))
    return out


def write_calderby(path, rows, skus):
    if os.path.exists(path):
        os.remove(path)
    con = sqlite3.connect(path)
    con.execute("PRAGMA page_size = 4096")
    con.execute("PRAGMA journal_mode = OFF")
    cur = con.cursor()
    cur.execute("CREATE TABLE stores (store_id INTEGER PRIMARY KEY, store_name TEXT NOT NULL, opened TEXT)")
    cur.execute("CREATE TABLE sku_gtin (sku TEXT PRIMARY KEY, gtin TEXT NOT NULL, description TEXT, first_sold TEXT)")
    cur.execute("CREATE TABLE daily_sales (sale_date TEXT NOT NULL, store_id INTEGER NOT NULL, sku TEXT NOT NULL, "
                "units INTEGER NOT NULL, net_value REAL NOT NULL)")
    cur.executemany("INSERT INTO stores VALUES (?,?,?)", [(s, n, o) for (s, n, _), o in
                                                       zip(W.CALDERBY_STORES, ("1998-04-06", "2007-11-19", "2016-06-27"))])
    cur.executemany("INSERT INTO sku_gtin VALUES (?,?,?,?)", skus)
    cur.executemany("INSERT INTO daily_sales VALUES (?,?,?,?,?)", [(a, b, c, d, float(e)) for a, b, c, d, e in rows])
    cur.execute("CREATE INDEX ix_sales_date ON daily_sales (sale_date)")
    con.commit()
    con.close()
    return len(rows)


def published_workbook(path, levels, base):
    import openpyxl
    from openpyxl.styles import Alignment, Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Item indices"
    months = ["%04d-%02d" % ym for ym in PB.PMONTHS]
    ws.append(["Kellesay Retail Prices Index: item indices (January 2015 = 100)"])
    ws["A1"].font = Font(bold=True)
    ws.append(["Published 15 October 2026. Indices to one decimal place; 12-month rates on the next sheet are "
               "calculated from the unrounded indices."])
    ws.append([])
    ws.append(["Item code", "Item"] + [dt.date(int(m[:4]), int(m[5:]), 1).strftime("%b %Y") for m in months])
    for code in W.ICODE:
        f = base[code] / levels[code][months[0]]
        ws.append([code, W.INAME[code]] + [round(levels[code][m] * f, 1) for m in months])
    ws.freeze_panes = "C5"
    ws.column_dimensions["B"].width = 26
    rs = wb.create_sheet("12-month rates")
    rmonths = months[12:]
    rs.append(["Kellesay Retail Prices Index: 12-month rates (per cent)"])
    rs["A1"].font = Font(bold=True)
    rs.append([])
    rs.append(["Item code", "Item"] + [dt.date(int(m[:4]), int(m[5:]), 1).strftime("%b %Y") for m in rmonths])
    for code in W.ICODE:
        row = [code, W.INAME[code]]
        for m in rmonths:
            prev = "%04d-%s" % (int(m[:4]) - 1, m[5:])
            row.append(round((levels[code][m] / levels[code][prev] - 1) * 100, 1))
        rs.append(row)
    rs.freeze_panes = "C4"
    rs.column_dimensions["B"].width = 26
    for sh in (ws, rs):
        for row in sh.iter_rows(min_row=4 if sh is ws else 3):
            for c in row[2:]:
                c.number_format = "0.0"
                c.alignment = Alignment(horizontal="right")
    wb.save(path)


def published_rates(levels):
    return {code: round((levels[code]["2026-09"] / levels[code]["2025-09"] - 1) * 100, 1) for code in W.ICODE}


def quote_rows(qrows):
    names = {o[0]: o[1] for o in PB.OUTLETS}
    out = []
    for q in qrows:
        if q["month"] < "2025-09":
            continue
        out.append([q["month"], q["item"], q["outlet"], names[q["outlet"]], q["pcode"], q["desc"], "%.2f" % q["price"],
                    q["status"], q["collected"]])
    out.sort(key=lambda r: (r[0], r[1], r[2], r[4]))
    return out


def returns_workbook(path, cv):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Wrenfold Stores"
    ws.append(["Wrenfold Stores Ltd: annual return of retail sales, calendar year 2025"])
    ws["A1"].font = Font(bold=True)
    ws.append(["Sales value including GST, pounds. Categories as in the 2025 product hierarchy."])
    ws.append([])
    ws.append(["Category code", "Category", "Sales 2025"])
    for r in cv["wren"]:
        ws.append(r)
    gs = wb.create_sheet("Gorsedale Co-operative")
    gs.append(["Gorsedale Co-operative Society: retail sales return 2025"])
    gs["A1"].font = Font(bold=True)
    gs.append(["Sales value including GST, pounds, by channel. Trade: sales on trade accounts."])
    gs.append([])
    gs.append(["Channel", "Category code", "Category", "Sales 2025"])
    for r in cv["gors"]:
        gs.append(r)
    cs = wb.create_sheet("Calderby's")
    cs.append(["Calderby's: sales by department, January to December 2025"])
    cs["A1"].font = Font(bold=True)
    cs.append(["Values exclude GST."])
    cs.append([])
    cs.append(["Dept", "Department", "Net sales"])
    for r in cv["cald"]:
        cs.append(r)
    for sh in (ws, gs, cs):
        for row in sh.iter_rows(min_row=5):
            row[-1].number_format = "#,##0.00"
        sh.column_dimensions["B"].width = 30
        sh.column_dimensions["C"].width = 30
    wb.save(path)


def spending_workbook(path, cv):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Household spending 2025"
    ws.append(["Kellesay households' spending by RPI item, 2025 (thousand pounds, including GST)"])
    ws["A1"].font = Font(bold=True)
    ws.append(["Source: Household Expenditure Survey 2025 grossed to the island's private households, with the 2025 "
               "retail sales inquiry. Resident private households only."])
    ws.append([])
    ws.append(["Item code", "Item", "Spending 2025"])
    for code in W.ICODE:
        ws.append([code, W.INAME[code], cv["spend"][code]])
    for row in ws.iter_rows(min_row=5):
        row[2].number_format = "#,##0"
    ws.column_dimensions["B"].width = 28
    wb.save(path)


def rsi_workbook(path, seed):
    """The island's retail sales index (value and volume, 2022 = 100), a neighbouring statistic the office publishes."""
    import openpyxl
    from openpyxl.styles import Font
    rng = np.random.default_rng([seed, 401])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Retail sales index"
    ws.append(["Kellesay retail sales index, 2022 = 100, not seasonally adjusted"])
    ws["A1"].font = Font(bold=True)
    ws.append([])
    ws.append(["Month", "All retailing value", "All retailing volume", "Food stores value", "Food stores volume",
               "Non-food value", "Non-food volume"])
    season = [0.86, 0.84, 0.93, 0.96, 1.0, 1.02, 1.07, 1.10, 1.0, 0.98, 1.03, 1.21]
    lv = [104.2, 98.7, 106.8, 101.2, 102.1, 96.9]
    gr = [0.0042, 0.0006, 0.0045, 0.0009, 0.0037, 0.0002]
    d = dt.date(2025, 1, 1)
    k = 0
    while d <= dt.date(2026, 8, 1):
        s = season[d.month - 1]
        row = [d.strftime("%b %Y")]
        for i in range(6):
            row.append(round(lv[i] * (1 + gr[i]) ** k * s * float(np.exp(rng.normal(0, 0.012))), 1))
        ws.append(row)
        k += 1
        d = dt.date(d.year + (d.month == 12), d.month % 12 + 1, 1)
    ws.column_dimensions["A"].width = 12
    wb.save(path)
