"""The pack's documents: Statistics Kellesay's plan for introducing supermarket transaction data into the RPI, the
folder's data dictionary and its manifest."""
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import A4                                                # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import cm                                                    # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import pack as P    # noqa: E402
import world as W   # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("LiberationSerif", "Regular"), ("LiberationSerif-Bold", "Bold"), ("LiberationSans", "Regular"),
              ("LiberationSans-Bold", "Bold")):
    try:
        pdfmetrics.getFont(nm)
    except KeyError:
        pdfmetrics.registerFont(TTFont(nm, FONT_DIR + nm.split("-")[0] + "-" + f + ".ttf"))
INK = colors.HexColor("#1b1d1f")
GREY = colors.HexColor("#5b6067")
TEAL = colors.HexColor("#0f4c5c")
BODY = ParagraphStyle("body", fontName="LiberationSerif", fontSize=10.5, leading=14.5, textColor=INK, spaceAfter=6)
NUM = ParagraphStyle("num", parent=BODY, leftIndent=30, firstLineIndent=-30)
H1 = ParagraphStyle("h1", fontName="LiberationSans-Bold", fontSize=14.5, leading=18, textColor=TEAL, spaceAfter=2)
H1S = ParagraphStyle("h1s", fontName="LiberationSans", fontSize=9.5, leading=12.5, textColor=GREY, spaceAfter=10)
H2 = ParagraphStyle("h2", fontName="LiberationSans-Bold", fontSize=10.5, leading=14, textColor=INK, spaceBefore=7, spaceAfter=3)
CELL = ParagraphStyle("cell", fontName="LiberationSans", fontSize=8.8, leading=11, textColor=INK)
CELLB = ParagraphStyle("cellb", parent=CELL, fontName="LiberationSans-Bold")
SMALL = ParagraphStyle("small", fontName="LiberationSans", fontSize=8.5, leading=11, textColor=GREY)


def _doc(path, title, author, footer):
    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFont("LiberationSans", 7.5)
        canvas.setFillColor(GREY)
        canvas.drawString(2.2 * cm, 1.4 * cm, footer)
        canvas.drawRightString(A4[0] - 2.2 * cm, 1.4 * cm, "Page %d" % doc.page)
        canvas.restoreState()
    d = SimpleDocTemplate(path, pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=2.0 * cm,
                          bottomMargin=2.2 * cm, title=title, author=author, subject="", creator="", producer="")
    return d, frame


def _table(rows, widths):
    data = [[Paragraph(str(c), CELLB if i == 0 else CELL) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa0a6")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6eef0")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 4),
                           ("RIGHTPADDING", (0, 0), (-1, -1), 4)]))
    return t


def write_plan(path):
    doc, frame = _doc(path, "RPI transaction data: introduction plan", W.ORG,
                      "%s  |  Retail Prices Index: introducing supermarket transaction data  |  Version 2, 24 June 2026"
                      % W.ORG)
    s = [Paragraph("Retail Prices Index: Introducing Supermarket Transaction Data", H1),
         Paragraph("%s. Introduction plan, version 2, approved by the Chief Statistician on 24 June 2026. Owner: %s, %s."
                   % (W.ORG, W.PEOPLE["prices"], W.TITLE["prices"]), H1S)]
    P_ = lambda n, t: s.append(Paragraph("<b>%s</b>&nbsp;&nbsp;%s" % (n, t), NUM))
    s.append(Paragraph("1. Purpose", H2))
    P_("1.1", "From the January 2027 index the office intends to price the RPI's packaged grocery and household items "
              "from the sales records of the island's three supermarket groups, Wrenfold Stores, Gorsedale Co-operative "
              "and Calderby's, in place of the monthly field collection, item by item. This plan sets out the items, the "
              "parallel year on which each is assessed, the definitions used to put the supermarkets' records on one "
              "footing, and the test an item must pass before it moves.")
    s.append(Paragraph("2. Items and records", H2))
    P_("2.1", "The sixteen items in Table 1. Every other RPI item stays on its present sources in 2027.")
    s.append(_table([["Item code", "Item", "Item code", "Item"]] +
                    [[W.ITEMS[i][0], W.ITEMS[i][1], W.ITEMS[i + 8][0], W.ITEMS[i + 8][1]] for i in range(8)],
                    [2.0 * cm, 5.6 * cm, 2.0 * cm, 5.6 * cm]))
    s.append(Paragraph("Table 1. Items assessed for transaction data.", SMALL))
    P_("2.2", "Under the data supply agreements signed in 2025 each supermarket sends its full sales records: Wrenfold "
              "Stores weekly, Gorsedale Co-operative by fiscal period, Calderby's monthly as an extract of its "
              "point-of-sale database. Records for September 2025 to September 2026 form the parallel year, over which "
              "the field collection continued unchanged.")
    s.append(Paragraph("3. Definitions", H2))
    P_("3.1", "<b>Product line.</b> A GTIN sold by one supermarket; the same GTIN at two supermarkets is two lines. A line "
              "belongs to the RPI item the office's product classification gives it.")
    P_("3.2", "<b>Month.</b> A day belongs to its calendar month. A supermarket's week belongs to the month in which the "
              "week ends; Gorsedale Co-operative's fiscal calendar gives the dates of its weeks.")
    P_("3.3", "<b>Monthly price.</b> A line's price in a month is its sales value including GST divided by its units sold "
              "in the month, returns netted.")
    P_("3.4", "<b>Lines compared.</b> The assessment of the parallel year covers the lines each supermarket sold in every "
              "month from September 2025 to September 2026.")
    P_("3.5", "<b>Index.</b> Each item's index from transaction data is compiled monthly over the lines compared, "
              "September 2025 = 100; a line counts in its item's price change by its share of the item's spending, as in "
              "a Törnqvist index.")
    P_("3.6", "<b>12-month rate.</b> An item's 12-month rate from transaction data is the percentage change in its index "
              "from September 2025 to September 2026, the change in the item's prices over the year.")
    s.append(Paragraph("4. The test", H2))
    P_("4.1", "An item moves to transaction data from the January 2027 index when its 12-month rate from transaction data "
              "is within 1.5 percentage points of its published 12-month rate to September 2026, either way. An item "
              "further apart stays on the field collection for 2027 while the office examines the difference.")
    P_("4.2", "The notice of the January 2027 changes is published on 13 November 2026, two months before the index is "
              "compiled, as the Statistics Code requires.")
    s.append(Spacer(1, 10))
    s.append(Paragraph("Supersedes version 1 (3 February 2026), which proposed a six-month parallel run.", SMALL))
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


def write_dictionary(path, counts):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    F = W.FILES
    c = lambda k: format(counts[k], ",")
    rows = [
        ("(all files)", "", "Dates are Kellesay dates (the island keeps British time) unless a column says otherwise. Money "
                            "in pounds. GST is charged at 5 per cent on every item in this folder."),
        (F["wrenfold"], "", "Wrenfold Stores' weekly sales by product, all stores, weeks Sunday to Saturday. %s rows." %
         c("wrenfold")),
        (F["wrenfold"], "week_ending", "The Saturday ending the week."),
        (F["wrenfold"], "units, sales_value", "Units sold, net of returns, and their value including GST."),
        (F["wrenfold"], "extract_date", "The date the week's records were extracted from Wrenfold's systems."),
        (F["gorsedale"], "", "Gorsedale Co-operative's sales by fiscal year, period and week, all stores, pipe-delimited. "
                             "%s rows." % c("gorsedale")),
        (F["gorsedale"], "WEEK", "Week of the fiscal year; weeks run Monday to Sunday."),
        (F["gorsedale"], "ITEM_NO, EAN", "Gorsedale's item number and the product's barcode number."),
        (F["gorsedale"], "QTY, SALES", "Units sold, net of returns, and their value including GST."),
        (F["calendar"], "", "Gorsedale Co-operative's fiscal calendar: each week's period and dates."),
        (F["calderby"], "", "Extract of Calderby's point-of-sale database: tables daily_sales (one row per store, day and "
                            "SKU, %s rows), sku_gtin and stores." % c("calderby")),
        (F["calderby"], "daily_sales.units, net_value", "Units sold (a return is a row of negative units) and their value "
                                                         "net of GST."),
        (F["calderby"], "sku_gtin", "Calderby's SKU, the product's GTIN and the date the SKU was first sold."),
        (F["classes"], "", "The office's product classification: each supermarket's product lines and the RPI item each "
                           "belongs to, with the dates the assignment holds. %s rows." % c("classes")),
        (F["published"], "", "The published RPI item indices, January 2024 to September 2026, and the items' 12-month "
                             "rates."),
        (F["deliveries"], "", "Every file the office received from the supermarkets under the data supply agreements. %s "
                              "rows." % c("deliveries")),
        (F["deliveries"], "covers_from, covers_to", "The first and last day of the records in the file."),
        (F["deliveries"], "received_utc", "When the file reached the office's server, UTC."),
        (F["quotes"], "", "The field collection's price quotes for the sixteen items, September 2025 to September 2026: the "
                          "collector's shelf price for each sampled product at each outlet. %s rows." % c("quotes")),
        (F["quotes"], "status", "C collected; T not available on the day, the last price carried."),
        (F["returns"], "", "Each supermarket's annual return of 2025 retail sales to the office, one sheet each, as "
                           "received."),
        (F["catmap"], "", "The office's map of the supermarkets' product categories to RPI items."),
        (F["spending"], "", "The office's estimate of Kellesay households' 2025 spending on each item."),
        (F["rsi"], "", "The office's monthly retail sales index, value and volume."),
        (F["manifest"], "", "The programme's register of the files in this folder: owner, period covered, version date, "
                            "records and handling."),
    ]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Field definitions"
    ws.append(["File", "Column", "What it holds"])
    for r in rows:
        ws.append([v or None for v in r])
    head = PatternFill("solid", fgColor="E6EEF0")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = head
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 90
    ws.freeze_panes = "A2"
    wb.save(path)


def write_manifest(path, counts):
    """The programme's register of the folder: each file's owner, what it covers, its version date, its records and how it
    is handled."""
    F = W.FILES
    sup, own, pub = "Supplier confidential", "Internal", "Published"
    prog, prices = "Transaction Data Programme", "Prices Statistics"
    rows = [
        [F["wrenfold"], "Wrenfold Stores", "Weekly item sales", "Weeks ending 6 Sep 2025 to 26 Sep 2026", "2026-10-01",
         counts["wrenfold"], sup],
        [F["gorsedale"], "Gorsedale Co-operative", "Item sales by fiscal week", "FY2026 P08 to FY2027 P08", "2026-10-06",
         counts["gorsedale"], sup],
        [F["calendar"], "Gorsedale Co-operative", "Fiscal calendar", "FY2026 and FY2027", "2026-02-02", counts["calendar"],
         sup],
        [F["calderby"], "Calderby's", "Point-of-sale database extract", "1 Sep 2025 to 30 Sep 2026", "2026-10-07",
         counts["calderby"], sup],
        [F["classes"], prices, "Supermarket lines classified to RPI items", "Lines sold from Sep 2025", "2026-10-14",
         counts["classes"], own],
        [F["published"], prices, "RPI item indices and 12-month rates", "Jan 2024 to Sep 2026", "2026-10-15", "", pub],
        [F["plan"], prog, "Introduction plan, version 2", "", "2026-06-24", "", own],
        [F["deliveries"], prog, "Data receipt log", "Files received to 21 Oct 2026", W.PACK_DATE.isoformat(),
         counts["deliveries"], own],
        [F["quotes"], prices, "RPI field collection quotes", "Sep 2025 to Sep 2026", "2026-10-09", counts["quotes"], own],
        [F["returns"], "Business Statistics", "Supermarkets' annual returns of retail sales", "2025", "2026-04-30", "",
         sup],
        [F["catmap"], prices, "Retailer category map", "", "2025-03-14", counts["catmap"], own],
        [F["spending"], "Household Statistics", "Household spending estimates by RPI item", "2025", "2026-07-22", "", own],
        [F["rsi"], "Business Statistics", "Retail sales index", "Jan 2025 to Aug 2026", "2026-09-18", "", pub],
        [F["dict"], prog, "Field definitions for this folder", "", W.PACK_DATE.isoformat(), "", own],
    ]
    return P.write_csv(path, ["file", "owner", "description", "period_covered", "version_date", "records", "handling"],
                       rows)
