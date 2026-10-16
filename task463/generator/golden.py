"""Compute the golden from the shipped pack and write the deliverables: the readiness workbook, the conformed
product-month file and the chart of each item's two 12-month rates; and generator/figures.json, every figure the
submission states."""
import csv
import json
import os
import sys

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import bundle as B     # noqa: E402
import methods as M    # noqa: E402
import world as W      # noqa: E402


def r1(x):
    """One decimal, half away from zero (the office's rounding)."""
    return float(np.sign(x) * np.floor(abs(x) * 10 + 0.5) / 10)


def compute(target):
    pm = B.product_months(target)
    cl = B.read_classes(target)
    arr = B.item_arrays(target, pm, cl)
    pub = B.published_rates(target)
    F = {"items": W.ICODE, "names": W.INAME, "line": W.LINE, "n_items": len(W.ICODE), "rate_months": 12,
         "first_index": "January 2027"}
    rate, diff, lines, move, chained, dfree, monthly, drift = {}, {}, {}, {}, {}, {}, {}, {}
    for it in W.ICODE:
        keys, P, Q, E = arr[it]
        g = M.rate(M.tornqvist(P, Q, E, 0, W.NM - 1))
        rate[it] = r1(g)
        diff[it] = r1(rate[it] - pub[it])
        lines[it] = len(keys)
        move[it] = abs(rate[it] - pub[it]) <= W.LINE + 1e-9
        chained[it] = r1(M.rate(M.chained(M.tornqvist, P, Q, E)))
        drift[it] = r1(g - M.rate(M.chained(M.tornqvist, P, Q, E)))
        dfree[it] = {k: round(v, 4) for k, v in M.drift_free(P, Q, E).items()}
        monthly[it] = [100.0] + [round(100.0 * M.tornqvist(P, Q, E, 0, m), 2) for m in range(1, W.NM)]
    lv = B.published_levels(target)
    F.update(rate=rate, published=pub, diff=diff, lines=lines, move={k: bool(v) for k, v in move.items()},
             chained=chained, linked_drift=drift, drift_free=dfree, monthly_index=monthly,
             published_monthly={it: [round(100.0 * lv[it][m] / lv[it][B.MONTH_KEYS[0]], 2) for m in B.MONTH_KEYS]
                                for it in W.ICODE})
    F["moved"] = [it for it in W.ICODE if move[it]]
    F["stay"] = [it for it in W.ICODE if not move[it]]
    F["n_moved"] = len(F["moved"])
    F["stay_gap"] = {it: abs(diff[it]) for it in F["stay"]}
    F["default_moved"] = [it for it in W.ICODE if abs(chained[it] - pub[it]) <= W.LINE + 1e-9]
    near_in = max(F["moved"], key=lambda it: (abs(diff[it]), it))
    near_out = min(F["stay"], key=lambda it: (abs(diff[it]), it))
    F["nearest_in"], F["nearest_out"] = near_in, near_out
    # the conformed product-months of the sixteen items
    rows = []
    panel = {(it, k) for it in W.ICODE for k in arr[it][0]}
    for (r, g, m), (u, v) in sorted(pm.items()):
        it = B.item_of(cl, r, g, m)
        if it not in W.INAME:
            continue
        rows.append([r, g, it, m, u, "%.2f" % v, ("%.4f" % (v / u)) if u > 0 else "", "Y" if (it, (r, g)) in panel else "N"])
    F["csv_rows"] = len(rows)
    F["csv_panel_lines"] = len(panel)
    F["csv_lines"] = len({(r[0], r[1]) for r in rows})
    # Ask A: each supermarket's average days from a month's end to its records being complete
    dl = B.delivery_figures(target)
    F["delivery"] = {r: r1(v["average"]) for r, v in dl.items()}
    F["delivery_lags"] = {r: v["lags"] for r, v in dl.items()}
    # Ask B: coverage of households' 2025 spending by the supermarkets' 2025 retail sales
    cov = B.coverage_figures(target)
    F["coverage"] = {it: int(np.floor(cov[it] + 0.5)) for it in W.ICODE}
    F["coverage_exact"] = {it: round(cov[it], 3) for it in W.ICODE}
    return F, rows


def write_workbook(path, F):
    import openpyxl
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Readiness 2027"
    bold = Font(bold=True)
    head = PatternFill("solid", fgColor="E6EEF0")
    thin = Side(style="thin", color="9AA0A6")
    ws.append(["RPI items moving to supermarket transaction data in January 2027: %d of 16" % F["n_moved"]])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Parallel year September 2025 to September 2026. Transaction-data 12-month rate against the published "
               "12-month rate; an item moves when the two are within 1.5 percentage points (introduction plan, 4.1)."])
    ws.append([])
    ws.append(["Item code", "Item", "Lines compared", "12-month rate, transaction data (%)",
               "12-month rate, published (%)", "Difference (points)", "January 2027"])
    for c in ws[4]:
        c.font = bold
        c.fill = head
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = Border(bottom=thin)
    for it in W.ICODE:
        ws.append([it, W.INAME[it], F["lines"][it], F["rate"][it], F["published"][it], F["diff"][it],
                   "Moves to transaction data" if F["move"][it] else "Stays on field collection"])
    for row in ws.iter_rows(min_row=5, max_row=20):
        row[2].number_format = "0"
        for c in row[3:6]:
            c.number_format = "0.0"
    ws.append([])
    ws.append(["Moving: %s." % ", ".join(W.INAME[it] for it in F["moved"])])
    ws.append(["Staying: %s." % ", ".join(W.INAME[it] for it in F["stay"])])
    ws.freeze_panes = "C5"
    for col, wdt in zip("ABCDEFG", (10, 26, 10, 16, 16, 13, 26)):
        ws.column_dimensions[col].width = wdt
    ws.row_dimensions[4].height = 44
    # Ask A
    ts = wb.create_sheet("Supplier timeliness")
    ts.append(["Days from a month's end until the supermarket's records for the month were complete"])
    ts["A1"].font = Font(bold=True, size=12)
    ts.append(["Average over the thirteen months of the parallel year; the arrival of the last file carrying the month's "
               "records, in Kellesay time. A week's records belong to the month in which the week ends."])
    ts.append([])
    ts.append(["Supermarket", "Average days"] + B.MONTH_KEYS)
    for c in ts[4]:
        c.font = bold
        c.fill = head
    for r in (W.RNAME[x] for x in W.RETAILERS):
        ts.append([r, F["delivery"][r]] + F["delivery_lags"][r])
    for row in ts.iter_rows(min_row=5, max_row=7):
        row[1].number_format = "0.0"
    ts.column_dimensions["A"].width = 26
    ts.column_dimensions["B"].width = 13
    # Ask B
    cs = wb.create_sheet("Coverage 2025")
    cs.append(["The three supermarkets' 2025 retail sales as a share of households' 2025 spending, by item"])
    cs["A1"].font = Font(bold=True, size=12)
    cs.append(["Retail sales with GST from the supermarkets' annual returns (Gorsedale's trade sales left out, "
               "Calderby's net sales with GST added), each category placed by its name in the office's category map."])
    cs.append([])
    cs.append(["Item code", "Item", "Coverage (%)"])
    for c in cs[4]:
        c.font = bold
        c.fill = head
    for it in W.ICODE:
        cs.append([it, W.INAME[it], F["coverage"][it]])
    cs.column_dimensions["B"].width = 26
    cs.column_dimensions["C"].width = 13
    # a chart sheet: the two rates side by side
    ch = BarChart()
    ch.type = "bar"
    ch.title = "12-month rate to September 2026 (%)"
    ch.y_axis.title = "per cent"
    data = Reference(ws, min_col=4, max_col=5, min_row=4, max_row=20)
    cats = Reference(ws, min_col=2, min_row=5, max_row=20)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ch.height, ch.width = 12, 18
    ws.add_chart(ch, "I4")
    ns = wb.create_sheet("Notes")
    for line in [
        "Product line: a GTIN at one supermarket. Lines compared: those each supermarket sold in every month from "
        "September 2025 to September 2026 (plan 3.4).",
        "Monthly price: sales value with GST over units sold in the month, returns netted; Calderby's net values carry "
        "5 per cent GST; each resent Wrenfold week at its latest extract; a retailer week in the month it ends.",
        "12-month rate: September 2026 compared with September 2025 directly over the lines compared, each line weighted "
        "by its spending share averaged over the two months (Törnqvist). Linking the twelve month-on-month comparisons "
        "is not used: offer months lift a stored good's sales and the month after falls, so the links drift below "
        "every line's own change.",
        "Published rates as published on 15 October 2026. Rates and differences to one decimal place; the test uses the "
        "rounded rates.",
    ]:
        ns.append([line])
        ns.cell(row=ns.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ns.column_dimensions["A"].width = 120
    wb.save(path)


def write_chart(path, F):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 7.5})
    fig, axes = plt.subplots(4, 4, figsize=(11.0, 9.6), dpi=110, sharex=True)
    x = list(range(W.NM))
    move_c, stay_c, pub_c = "#0f4c5c", "#b5462c", "#8a8f96"
    for ax, it in zip(axes.flat, W.ICODE):
        t, p = F["monthly_index"][it], F["published_monthly"][it]
        ax.axhline(100, color="#d5d8db", lw=0.7, zorder=0)
        ax.plot(x, p, color=pub_c, lw=1.2)
        ax.plot(x, t, color=move_c if F["move"][it] else stay_c, lw=1.7)
        ax.set_title("%s %s\n%s: %.1f against %.1f published" % (
            it, W.INAME[it], "moves" if F["move"][it] else "stays", F["rate"][it], F["published"][it]),
            fontsize=7.5, loc="left", linespacing=1.3, fontweight="normal" if F["move"][it] else "bold",
            color="#222222" if F["move"][it] else stay_c)
        ax.set_xticks([0, 4, 8, 12])
        ax.set_xticklabels(["Sep 25", "Jan 26", "May 26", "Sep 26"])
        ax.tick_params(length=2.5)
        for sp in ("top", "right"):
            ax.spines[sp].set_visible(False)
    handles = [Line2D([], [], color=move_c, lw=1.7, label="Transaction data, item moves"),
               Line2D([], [], color=stay_c, lw=1.7, label="Transaction data, item stays"),
               Line2D([], [], color=pub_c, lw=1.2, label="Published RPI item index")]
    fig.legend(handles=handles, loc="upper left", bbox_to_anchor=(0.01, 0.955), ncol=3, frameon=False, fontsize=8)
    fig.suptitle("%d of 16 items' transaction-data indices end within 1.5 points of the published RPI" % F["n_moved"],
                 x=0.01, y=0.995, ha="left", fontsize=11, fontweight="bold")
    fig.text(0.99, 0.005, "Indices September 2025 = 100; titles give each item's 12-month rate to September 2026, per cent. "
             "Statistics Kellesay, parallel year September 2025 to September 2026.", ha="right", fontsize=6.5,
             color="#555555")
    fig.tight_layout(rect=(0, 0.015, 1, 0.94), h_pad=1.6)
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)


def main(target, gdir, fig_path):
    os.makedirs(gdir, exist_ok=True)
    F, rows = compute(target)
    write_workbook(os.path.join(gdir, W.GOLDEN["workbook"]), F)
    with open(os.path.join(gdir, W.GOLDEN["products"]), "w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh, lineterminator="\r\n")
        wr.writerow(["retailer", "gtin", "rpi_item", "month", "units", "sales_value_incl_gst", "unit_price",
                     "in_parallel_year_lines"])
        wr.writerows(rows)
    write_chart(os.path.join(gdir, W.GOLDEN["chart"]), F)
    with open(fig_path, "w") as fh:
        json.dump(F, fh, indent=1, sort_keys=True)
        fh.write("\n")
    return F


if __name__ == "__main__":
    T = os.path.dirname(HERE)
    main(os.path.join(T, "target"), os.path.join(T, "golden"), os.path.join(HERE, "figures.json"))
