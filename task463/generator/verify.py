"""Independent verifier for task463: reads only the shipped folder and the goldens, on its own code path (pandas, the
sqlite3 module, zoneinfo), and recomputes each item's 12-month rate from transaction data, the moved items, each
supermarket's average days to complete records and each item's 2025 coverage, against generator/figures.json and the
golden workbook. Imports nothing from the generator.

    python3 task463/generator/verify.py
"""
import json
import math
import os
import sqlite3
import sys
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
TASK = os.path.dirname(HERE)
LINE = 1.5
MONTHS = pd.period_range("2025-09", "2026-09", freq="M").astype(str).tolist()
ITEMS = ["1101", "1103", "1105", "1108", "1145", "1146", "1151", "1182", "1211", "1212", "1221", "1222", "1223", "5611",
         "5613", "6213"]


def find(target, stem):
    hit = [f for f in os.listdir(target) if f.startswith(stem)]
    assert len(hit) == 1, (stem, hit)
    return os.path.join(target, hit[0])


def r1(x):
    return math.copysign(math.floor(abs(x) * 10 + 0.5) / 10, x)


def product_months(target):
    w = pd.read_csv(find(target, "wrenfold_weekly"), dtype={"gtin": str})
    w = w.sort_values("extract_date").groupby(["week_ending", "gtin"], as_index=False).last()
    w["month"] = pd.to_datetime(w["week_ending"]).dt.strftime("%Y-%m")
    w = w.assign(retailer="Wrenfold Stores").rename(columns={"sales_value": "value"})[["retailer", "gtin", "month", "units",
                                                                                         "value"]]
    cal = pd.read_csv(find(target, "gorsedale_fiscal_calendar"))
    g = pd.read_csv(find(target, "gorsedale_period_sales"), sep="|", dtype={"EAN": str})
    g = g.merge(cal, left_on=["FISCAL_YEAR", "WEEK"], right_on=["fiscal_year", "week"])
    g["month"] = pd.to_datetime(g["week_end"]).dt.strftime("%Y-%m")
    g["gtin"] = g["EAN"].str.zfill(13)
    g = g.assign(retailer="Gorsedale Co-operative").rename(columns={"QTY": "units", "SALES": "value"})[
        ["retailer", "gtin", "month", "units", "value"]]
    con = sqlite3.connect(find(target, "calderby_pos"))
    c = pd.read_sql("SELECT s.sale_date, k.gtin, s.units, s.net_value FROM daily_sales s JOIN sku_gtin k USING (sku)", con)
    con.close()
    c["month"] = c["sale_date"].str[:7]
    c["value"] = c["net_value"] * 1.05
    c = c.assign(retailer="Calderby's")[["retailer", "gtin", "month", "units", "value"]]
    pm = pd.concat([w, g, c]).groupby(["retailer", "gtin", "month"], as_index=False)[["units", "value"]].sum()
    pm = pm[pm["month"].isin(MONTHS)]
    cl = pd.read_csv(find(target, "statistics_kellesay_product_classification"), dtype={"gtin": str, "rpi_item": str})
    cl = cl[cl["valid_to"].isna()][["retailer", "gtin", "rpi_item"]]
    return pm.merge(cl, on=["retailer", "gtin"])


def item_rates(pm):
    out, lines = {}, {}
    for it in ITEMS:
        d = pm[(pm["rpi_item"] == it) & (pm["units"] > 0)]
        full = d.groupby(["retailer", "gtin"])["month"].nunique()
        keep = full[full == len(MONTHS)].index
        d = d.set_index(["retailer", "gtin"]).loc[keep].reset_index()
        a = d[d["month"] == "2025-09"].set_index(["retailer", "gtin"])
        b = d[d["month"] == "2026-09"].set_index(["retailer", "gtin"])
        pa, pb = a["value"] / a["units"], b["value"] / b["units"]
        sa, sb = a["value"] / a["value"].sum(), b["value"] / b["value"].sum()
        out[it] = r1((math.exp(float((0.5 * (sa + sb) * np.log(pb / pa)).sum())) - 1) * 100)
        lines[it] = len(keep)
    return out, lines


def published(target):
    x = pd.read_excel(find(target, "kellesay_rpi_item_indices"), sheet_name="12-month rates", header=2, dtype={"Item code": str})
    return dict(zip(x["Item code"], x["Sep 2026"]))


def delivery(target):
    d = pd.read_csv(find(target, "retailer_delivery_log"))
    d["local"] = pd.to_datetime(d["received_utc"]).dt.tz_convert(ZoneInfo("Europe/London")).dt.date
    res = {}
    for r, g in d.groupby("retailer"):
        done = {}
        for _, row in g.iterrows():
            a, b = pd.Timestamp(row["covers_from"]), pd.Timestamp(row["covers_to"])
            if r == "Calderby's":
                months = set(pd.date_range(a, b).strftime("%Y-%m"))
            else:
                months = set(pd.date_range(a + pd.Timedelta(days=6), b, freq="7D").strftime("%Y-%m"))
            for m in months & set(MONTHS):
                done[m] = max(done.get(m, row["local"]), row["local"])
        lags = [(done[m] - (pd.Period(m, "M").end_time.date())).days for m in MONTHS]
        res[r] = r1(float(np.mean(lags)))
    return res


def coverage(target):
    path = find(target, "retailer_annual_returns")
    cmap = pd.read_csv(find(target, "retailer_category_map"), dtype=str)
    names = dict(zip(zip(cmap["retailer"], cmap["category_name"]), cmap["rpi_item"]))
    sales = {}
    wr = pd.read_excel(path, sheet_name="Wrenfold Stores", header=3)
    for _, r in wr.iterrows():
        it = names.get(("Wrenfold Stores", r["Category"]))
        sales[it] = sales.get(it, 0) + r["Sales 2025"]
    gs = pd.read_excel(path, sheet_name="Gorsedale Co-operative", header=3)
    for _, r in gs[gs["Channel"] == "Retail"].iterrows():
        it = names.get(("Gorsedale Co-operative", r["Category"]))
        sales[it] = sales.get(it, 0) + r["Sales 2025"]
    cs = pd.read_excel(path, sheet_name="Calderby's", header=3)
    for _, r in cs.iterrows():
        it = names.get(("Calderby's", r["Department"]))
        sales[it] = sales.get(it, 0) + r["Net sales"] * 1.05
    sp = pd.read_excel(find(target, "household_spending_estimates"), header=3, dtype={"Item code": str})
    spend = dict(zip(sp["Item code"], sp["Spending 2025"]))
    return {it: int(math.floor(100 * sales[it] / 1000 / spend[it] + 0.5)) for it in ITEMS}


def main(target, golden, fig_path=None):
    F = json.load(open(fig_path or os.path.join(HERE, "figures.json")))
    pm = product_months(target)
    rates, lines = item_rates(pm)
    pub = published(target)
    mv = sorted(it for it in ITEMS if abs(rates[it] - pub[it]) <= LINE + 1e-9)
    bad = []

    def check(label, cond, detail=""):
        print("%s  verify: %s%s" % ("PASS" if cond else "FAIL", label, ("  [" + str(detail) + "]") if detail else ""))
        if not cond:
            bad.append(label)
    check("each item's 12-month rate from transaction data, one decimal", rates == F["rate"],
          {k: (rates[k], F["rate"][k]) for k in ITEMS if rates[k] != F["rate"][k]})
    check("each item's lines compared", lines == F["lines"])
    check("the items that move", mv == sorted(F["moved"]), ", ".join(mv))
    check("the published rates", {k: float(v) for k, v in pub.items()} == {k: float(v) for k, v in F["published"].items()})
    dv = delivery(target)
    check("each supermarket's average days to complete records", dv == F["delivery"], dv)
    cv = coverage(target)
    check("each item's 2025 coverage", cv == F["coverage"], {k: (cv[k], F["coverage"][k]) for k in ITEMS if cv[k] != F["coverage"][k]})
    import openpyxl
    ws = openpyxl.load_workbook(os.path.join(golden, "kellesay_rpi_transaction_readiness_2027.xlsx"))["Readiness 2027"]
    rows = {str(ws.cell(row=r, column=1).value): [ws.cell(row=r, column=c).value for c in range(3, 8)] for r in range(5, 21)}
    check("the golden workbook's rows", all(rows[it][0] == lines[it] and abs(rows[it][1] - rates[it]) < 1e-9 and
                                            abs(rows[it][2] - pub[it]) < 1e-9 and
                                            rows[it][4].startswith("Moves") == (it in mv) for it in ITEMS))
    if bad:
        raise SystemExit("verify failed: %s" % bad)
    return rates


if __name__ == "__main__":
    main(os.path.join(TASK, "target"), os.path.join(TASK, "golden"))
