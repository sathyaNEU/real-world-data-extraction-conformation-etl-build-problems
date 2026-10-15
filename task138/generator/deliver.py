"""The golden deliverables: kbid_order_2027.xlsx (the order, the thirty replays, the 2026 lateral balances and the
notes) and kbid_order_2027.svg (the thirty replays from smallest to largest with the order drawn across), plus
generator/figures.json. Every figure is computed from the shipped pack (calc.py)."""
from __future__ import annotations

import json
import math
import os

import numpy as np
import openpyxl
from openpyxl.styles import Border, Side, Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import calc as C
import common as K
import postprocess as PP

MONTHS = ["February", "March", "April", "May", "June", "July", "August", "September", "October"]


def share1(x):
    return math.floor(x * 10 + 0.5) / 10.0


def compute(target):
    M = C.load(target)
    lad = C.ladder(M)
    g = lad["G"]
    A = C.load_asks(target)
    deliv, share, inflow = C.ask_a(A)
    spill = C.ask_b(A)
    meth = np.array([f["method"] for f in M.fields])
    held = meth != "Border"
    return dict(M=M, lad=lad, g=g, A=A, deliv=deliv, share=share, inflow=inflow, spill=spill, held=held)


def figures(c):
    g, lad, M = c["g"], c["lad"], c["M"]
    out = {
        "call": {"order_af": "{:,}".format(g["order"])},
        "steps": {
            "fields": M.n, "acres": "{:,.0f}".format(M.acres.sum()),
            "held_fields": int(c["held"].sum()),
            # the 24th smallest replay lies inside the order's thousand (every converging construction too)
            "setting_replay_between_af": ["{:,}".format(g["order"] - K.ORDER_STEP), "{:,}".format(g["order"])],
            "setting_year": g["year"],
            "replays_covered": 24,
            "order_af": "{:,}".format(g["order"]),
            # the Water Order Procedure clauses the write-up's steps cite
            "procedure_clauses": ["2.3", "2.4", "3.1", "3.2", "3.3", "3.4", "3.5", "4.1", "6.1", "6.2"],
            "table": "4",
        },
        "stations": {"held_stations": sum(1 for st in M.stations if st["method"] != "Border"),
                     "stations": len(M.stations),
                     "above_table4": all(v > M.T["ea"][next(st["method"] for st in M.stations if st["station"] == k)]
                                         for k, v in M.ea_station.items())},
        "default_path": {"order_af": "{:,}".format(lad["T4"]["order"]), "setting_year": lad["T4"]["year"]},
        "rungs": {k: {"order_af": "{:,}".format(v["order"]), "setting_year": v["year"]} for k, v in lad.items()},
        "ask_deliveries_af": {l: "{:,}".format(K.round_to(c["deliv"][l], 10)) for l in K.LATERALS},
        "ask_water_in_af": {l: "{:,}".format(K.round_to(c["inflow"][l], 10)) for l in K.LATERALS},
        "ask_not_delivered_pct": {l: "%.1f" % share1(c["share"][l]) for l in K.LATERALS},
        "ask_tail_weir_af": {l: "{:,}".format(K.round_to(c["spill"][l], 10)) for l in K.LATERALS},
    }
    return out


def _head(ws, row, values, widths=None):
    for k, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=k, value=v)
        cell.font = Font(name="Calibri", bold=True)
        cell.fill = PatternFill("solid", fgColor="DCE6DC")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    if widths:
        for k, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(k)].width = w


def write_book(path, c, figs):
    g = c["g"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Order 2027"
    ws["A1"] = "Kestrel Bend Irrigation District orders %s acre-feet from the Tolan Valley Project for 2027" % (
        "{:,}".format(g["order"]))
    ws["A1"].font = Font(name="Calibri", bold=True, size=13)
    ws["A2"] = ("Water Resources, under the Water Order Procedure (rev. December 2025) and its tables for the 2027 order. "
                "Prepared 1 February 2027.")
    ws["A2"].font = Font(name="Calibri", italic=True, size=10, color="595959")
    rows = [("2027 order (acre-feet)", g["order"]),
            ("Season", "1 February to 31 October 2027"),
            ("Replayed seasons the order covers", "24 of 30 (weather years 1997 to 2026)"),
            ("Replay that sets the order", g["year"])]
    r = 4
    for label, v in rows:
        ws.cell(row=r, column=1, value=label).font = Font(name="Calibri", bold=True)
        cell = ws.cell(row=r, column=2, value=v)
        if isinstance(v, int) and v > 3000:
            cell.number_format = "#,##0"
        r += 1
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 38
    r += 1
    ws.cell(row=r, column=1, value="Each replay starts every field from its root zone at the end of 31 January 2027, "
                                    "from its 2026 water at the share the soil moisture stations show reached the root "
                                    "zone, and the record's weather since (Notes).")
    ws.cell(row=r, column=1).font = Font(name="Calibri", italic=True, size=9, color="595959")

    ws2 = wb.create_sheet("Replays")
    ws2["A1"] = "The order covers 24 of the 30 replayed weather years; six would need more"
    ws2["A2"] = "Headgate requirement of the 2027 fields with each weather year's 1 February to 31 October weather (WOP 2.4)."
    ws2["A2"].font = Font(name="Calibri", italic=True, size=10, color="595959")
    ws2["A1"].font = Font(name="Calibri", bold=True, size=12)
    ranked = [y for y, _ in g["ranked"]]
    _head(ws2, 3, ["Weather year", "Headgate requirement (acre-feet)", "Rank (smallest = 1)", "Covered by the order"],
          [14, 30, 18, 20])
    rr = 3
    for y in sorted(g["reqs"]):
        rr += 1
        ws2.cell(row=rr, column=1, value=y)
        ws2.cell(row=rr, column=2, value=K.round_to(g["reqs"][y], 10)).number_format = "#,##0"
        ws2.cell(row=rr, column=3, value=ranked.index(y) + 1)
        ws2.cell(row=rr, column=4, value="yes" if g["reqs"][y] <= g["order"] else "no")
    ws2.freeze_panes = "A4"
    ws2.print_title_rows = "3:3"

    ws4 = wb.create_sheet("Laterals 2026")
    top = max(K.LATERALS, key=lambda l: c["share"][l])
    ws4["A1"] = "%s lost the largest share of its water in 2026: %.1f per cent against a %.1f per cent allowance" % (
        K.LATERAL_NAME[top], share1(c["share"][top]), round(K.LOSS[top] * 100, 1))
    ws4["A1"].font = Font(name="Calibri", bold=True, size=12)
    ws4["A2"] = ("From the riders' meter book with the meter shop's work orders, the heading loggers' daily totals with the "
                 "Water Master's estimates, and the riders' tail weir sheets.")
    ws4["A2"].font = Font(name="Calibri", italic=True, size=10, color="595959")
    _head(ws4, 3, ["Lateral", "Turnout deliveries (acre-feet)", "Water put in at the heading (acre-feet)",
                   "Share of the water put in not delivered (%)", "Over the tail weir (acre-feet)",
                   "Loss allowance, Table 5 (%)"], [12, 20, 22, 22, 18, 16])
    rr = 3
    for l in K.LATERALS:
        rr += 1
        ws4.cell(row=rr, column=1, value=K.LATERAL_NAME[l])
        ws4.cell(row=rr, column=2, value=K.round_to(c["deliv"][l], 10)).number_format = "#,##0"
        ws4.cell(row=rr, column=3, value=K.round_to(c["inflow"][l], 10)).number_format = "#,##0"
        ws4.cell(row=rr, column=4, value=share1(c["share"][l])).number_format = "0.0"
        ws4.cell(row=rr, column=5, value=K.round_to(c["spill"][l], 10)).number_format = "#,##0"
        ws4.cell(row=rr, column=6, value=round(K.LOSS[l] * 100, 1)).number_format = "0.0"
    rr += 1
    tin = sum(c["inflow"].values())
    tdl = sum(c["deliv"].values())
    for k, v in enumerate(["Nine laterals", K.round_to(tdl, 10), K.round_to(tin, 10), share1(100.0 * (tin - tdl) / tin),
                           K.round_to(sum(c["spill"].values()), 10), None], 1):
        cell = ws4.cell(row=rr, column=k, value=v)
        cell.font = Font(name="Calibri", bold=True)
        cell.border = Border(top=Side(style="thin", color="7F7F7F"))
        if k in (2, 3, 5):
            cell.number_format = "#,##0"
        elif k == 4:
            cell.number_format = "0.0"
    ws4.freeze_panes = "B4"
    ws4.page_setup.orientation = "landscape"

    ws5 = wb.create_sheet("Notes")
    ws5["A1"] = "Notes"
    ws5["A1"].font = Font(name="Calibri", bold=True, size=12)
    _head(ws5, 3, ["Item", "Note"], [24, 112])
    notes = [
        ("Order", "Smallest whole number of thousands of acre-feet that would have covered the 2027 headgate "
                  "requirement in at least 24 of the weather years 1997 to 2026 (WOP 2.3)."),
        ("Fields", "Latest accepted 2027 request for each field: crop, method and acres from the request; turnout, "
                   "lateral and soil unit from the register as it stood on 1 February 2027 (WOP 6.2). Withdrawn "
                   "requests and superseded versions not run."),
        ("Root zone at 31 January", "Each field run day by day through the WOP 3.1 to 3.3 balance from the end of 27 "
                                    "January 2026, when the winter's rains had every root zone at field capacity, to "
                                    "the end of 31 January 2027 on the record's weather, each delivered 2026 set at its "
                                    "ordered flow for the hours it ran, times the share of it that reached the root "
                                    "zone. That share is the one the soil moisture network's 2026 readings give for "
                                    "the field's method: each drip, micro-sprinkler and sprinkler station's season "
                                    "fitted (least squares) by the share its own sets reached its root zone at, "
                                    "averaged over the method's stations, above Table 4 at every station; border "
                                    "fields at Table 4, their sets running past field capacity. Cancelled orders carry "
                                    "no water."),
        ("Replays", "Each replay from that root zone, with WOP 3.4 and 3.5 as filed: an irrigation at the allowable "
                    "depletion returns the field to field capacity and delivers the depletion over the Table 4 "
                    "efficiency."),
        ("Weather record", "Station 39 to 30 April 2014 and the yard station from 1 May 2014; a day the yard station "
                           "did not record takes station 39's values (WOP 6.1)."),
        ("Laterals 2026", "Deliveries from the riders' meter reads, a meter changed in the season chained through the "
                          "shop's register readings at the change; the west spur turnouts (05-021 to 05-034) on "
                          "Lateral 5 to their 15 June reads and on Lateral 6 after. Water put in from the heading "
                          "loggers' daily totals, with the Water Master's estimates (operations log) for the days a "
                          "logger kept no total. Tail weir spill from the riders' morning gauge readings, each standing "
                          "until the next reading, the crest raised 2 inches for each board in place."),
        ("Rounding", "Order to the whole thousand acre-feet; other volumes to the nearest 10 acre-feet; shares to one "
                     "decimal place. Figures carried unrounded until shown."),
    ]
    rr = 3
    for item, note in notes:
        rr += 1
        ws5.cell(row=rr, column=1, value=item).alignment = Alignment(vertical="top")
        cell = ws5.cell(row=rr, column=2, value=note)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    PP.normalize_ooxml(path, K.PEOPLE["requester"], K.PEOPLE["requester"], "2027-01-28T17:40:12Z",
                       "2027-02-01T19:52:44Z", excel=True)


def write_chart(path, c, figs):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({"font.family": "DejaVu Sans", "svg.hashsalt": "task138-kbid", "svg.fonttype": "none"})
    g = c["g"]
    ranked = g["ranked"]
    years = [str(y) for y, _ in ranked]
    vals = [v / 1000.0 for _, v in ranked]
    covered = [v <= g["order"] for _, v in ranked]
    ink, muted, surface = "#1f2421", "#5b625d", "#fcfcfb"
    blue, orange = "#2a78d6", "#eb6834"
    fig, ax = plt.subplots(figsize=(10.5, 5.6), dpi=100)
    fig.patch.set_facecolor(surface)
    ax.set_facecolor(surface)
    x = np.arange(len(vals))
    cv = np.array(covered)
    ax.scatter(x[cv], np.array(vals)[cv], s=64, color=blue, edgecolors=surface, linewidths=1.2,
               label="Covered by the order", zorder=4)
    ax.scatter(x[~cv], np.array(vals)[~cv], s=64, color=orange, edgecolors=surface, linewidths=1.2,
               label="Above the order", zorder=4)
    order_k = g["order"] / 1000.0
    ax.axhline(order_k, color=ink, linewidth=1.4, zorder=3)
    ax.text(0.2, order_k + 0.6, "Order %s acre-feet" % "{:,}".format(g["order"]), color=ink, fontsize=10,
            va="bottom", ha="left", zorder=5)
    k24 = 23
    ax.annotate("%d replay sets the order" % ranked[k24][0],
                xy=(k24, vals[k24]), xytext=(k24 - 9.5, vals[k24] + 5.5), color=muted, fontsize=9,
                arrowprops=dict(arrowstyle="-", color=muted, linewidth=0.8))
    lo = math.floor(min(vals) / 10.0) * 10
    hi = math.ceil(max(vals) / 10.0) * 10 + 5
    ax.set_ylim(lo, hi)
    ax.set_xticks(x)
    ax.set_xticklabels(years, rotation=90, fontsize=8.5, color=muted)
    ax.tick_params(axis="y", colors=muted, labelsize=9)
    ax.set_ylabel("Headgate requirement, thousand acre-feet", color=muted, fontsize=10)
    ax.set_xlabel("Weather year replayed (smallest requirement to largest)", color=muted, fontsize=10)
    ax.grid(axis="y", color="#d9dcd9", linewidth=0.6, zorder=0)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    for s in ("left", "bottom"):
        ax.spines[s].set_color("#9ea79f")
    ax.legend(loc="upper left", frameon=False, fontsize=9, labelcolor=ink)
    ax.set_title("2027 order: %s acre-feet covers 24 of the 30 replayed seasons\n" % "{:,}".format(g["order"]),
                 loc="left", fontsize=13, color=ink, fontweight="bold")
    ax.text(0.0, 1.015, "Kestrel Bend Irrigation District, 2027 fields, each replay from the root zone on 31 January "
            "2027", transform=ax.transAxes, fontsize=9.5, color=muted, va="bottom", ha="left")
    fig.tight_layout()
    fig.savefig(path, facecolor=surface, metadata={"Date": None, "Creator": None, "Format": None, "Type": None})
    plt.close(fig)
    import re
    svg = open(path, encoding="utf-8").read()
    svg = re.sub(r"<!-- Created with matplotlib.*?-->\n?", "", svg)
    svg = re.sub(r"\s*<metadata>.*?</metadata>", "", svg, flags=re.S)
    svg = svg.replace('id="matplotlib.axis_', 'id="axis_')
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(svg)


def main(target, golden, figures_path):
    os.makedirs(golden, exist_ok=True)
    c = compute(target)
    figs = figures(c)
    write_book(os.path.join(golden, K.GOLD_BOOK), c, figs)
    write_chart(os.path.join(golden, K.GOLD_CHART), c, figs)
    with open(figures_path, "w") as fh:
        json.dump(figs, fh, indent=2, sort_keys=False)
        fh.write("\n")
    return c, figs
