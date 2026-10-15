"""The Water Order Procedure's arithmetic, every rung of the ladder and the four lateral asks, on inputs read from the
pack (load, load_asks). Vectorised over fields with numpy; verify.py recomputes the graded figures on its own code
path."""
from __future__ import annotations

import csv
import datetime as dt
import os

import numpy as np

import common as K


class Model:
    """The 2027 fields as the procedure sees them, the District's weather record and the procedure's tables."""

    def __init__(self, fields, wx, tables, last_delivery=None, turnout_last=None, cancelled_last=None, record=None):
        self.fields = fields
        self.ids = [f["id"] for f in fields]
        self.n = len(fields)
        self.wx = wx
        self.T = tables
        lat_index = {l: i for i, l in enumerate(K.LATERALS)}
        self.lat = np.array([lat_index[f["lateral"]] for f in fields])
        self.acres = np.array([f["acres"] for f in fields], dtype=float)
        self.kc = np.array([tables["kc"][f["crop"]] for f in fields], dtype=float)          # n x 12
        depth = []
        for f in fields:
            lim = tables["soils"][f["soil"]][1]
            root = tables["root"][f["crop"]]
            depth.append(root if lim is None else min(root, lim))
        self.taw = np.array([tables["soils"][f["soil"]][0] * d for f, d in zip(fields, depth)], dtype=float)
        self.raw = np.array([tables["mad"][f["crop"]] for f in fields], dtype=float) * self.taw
        self.ea = np.array([tables["ea"][f["method"]] for f in fields], dtype=float)
        self.loss = np.array([tables["loss"][l] for l in K.LATERALS], dtype=float)
        self.main = tables["main_loss"]
        self.last_delivery = last_delivery
        self.turnout_last = turnout_last
        self.cancelled_last = cancelled_last
        # the 2026 record: (field id, set start, set end, cfs, status) for every 2026 order on a 2027 field
        self.record = record or []
        self.index = {f: i for i, f in enumerate(self.ids)}
        self.stations = []

    # -------------------------------------------------------------------------------------------- the daily balance
    def day(self, D, d, cap=True):
        e, r = self.wx[d]
        D = D + self.kc[:, d.month - 1] * e - r
        D = np.maximum(D, 0.0)
        return np.minimum(D, self.taw) if cap else D

    def carry(self, anchors, upto=dt.date(2027, 1, 31), cap=True):
        """Each field's depletion at the end of `upto`, from zero at the end of its anchor day (an irrigation)."""
        D = np.zeros(self.n)
        a_ord = np.array([a.toordinal() for a in anchors])
        for d in K.daterange(min(anchors) + dt.timedelta(days=1), upto):
            nd = self.day(D, d, cap)
            D = np.where(d.toordinal() > a_ord, nd, D)
        return D

    def replay(self, year, D0, monthly=False, drop_leap=False):
        """WOP 2.4, 3.1 to 3.5: the 2027 fields under year's 1 February to 31 October weather from depletion D0.
        Returns the headgate requirement (AF) and, if asked, by month (February to October)."""
        D = D0.copy()
        mon = np.zeros((9, self.n))
        for d in K.season_days(year):
            if drop_leap and d.month == 2 and d.day == 29:
                continue
            D = self.day(D, d)
            irr = D >= self.raw
            mon[d.month - 2] += np.where(irr, D, 0.0)
            D = np.where(irr, 0.0, D)
        farm_m = mon / self.ea * self.acres / 12.0                      # AF at the turnouts, 9 x n
        by_lat = np.zeros((9, 9))
        for k in range(9):
            by_lat[k] = np.bincount(self.lat, weights=farm_m[k], minlength=9)
        head_m = (by_lat / (1 - self.loss)).sum(axis=1) / (1 - self.main)
        if monthly:
            return float(head_m.sum()), head_m
        return float(head_m.sum())

    def thirty(self, start, drop_leap=False):
        """start(year) -> D0. The thirty replays' requirements and monthly requirements."""
        req, mon = {}, {}
        for y in K.REPLAY_YEARS:
            tot, m = self.replay(y, start(y), monthly=True, drop_leap=drop_leap)
            req[y], mon[y] = tot, m
        return req, mon

    # -------------------------------------------------------------------------------------------- starts (rungs)
    def start_full(self):
        z = np.zeros(self.n)
        return lambda y: z

    def start_last_set(self):
        """R3: each field's last delivered 2026 set taken as returning its root zone to field capacity (zero at the end
        of the set's day), then the record's weather to 31 January 2027."""
        D = self.carry([self.last_delivery[i] for i in self.ids])
        self.D_last_set = D
        return lambda y: D

    # -------------------------------------------------------------------------------------------- the 2026 record
    def set_water(self, statuses=("Delivered",), split="hourly", acres=None):
        """{date: AF reaching each field that day} from the 2026 orders. split: hourly (the ordered flow for the
        hours the set ran each day), start (all on the set's first day) or end (all on its last day)."""
        idx = {f: i for i, f in enumerate(self.ids)}
        out = {}
        for fid, a, b, cfs, status in self.record:
            if status not in statuses or fid not in idx:
                continue
            i = idx[fid]
            if split == "hourly":
                t = a
                while t < b:
                    seg = min(b, dt.datetime(t.year, t.month, t.day) + dt.timedelta(days=1))
                    out.setdefault(t.date(), np.zeros(self.n))[i] += cfs * (seg - t).total_seconds() / K.CUFT_PER_AF
                    t = seg
            else:
                day = a.date() if split == "start" else b.date()
                out.setdefault(day, np.zeros(self.n))[i] += cfs * (b - a).total_seconds() / K.CUFT_PER_AF
        return out

    def hindcast(self, water, start=K.WINTER_FULL + dt.timedelta(days=1), D0=None, use_ea=True, acres=None,
                 order="net", cap=True, upto=None, ea=None, kcf=1.0):
        """Each field's root zone run day by day from `start` (field capacity unless D0) to the end of 31 January 2027
        through the procedure's daily balance on the record's weather, each day's irrigation water reaching the root
        zone at the share `ea` (by field; Table 4 when not given, all of it when use_ea is False). order: net (the
        day's ET, rain and irrigation together, then the floor and cap), after (the irrigation after the day's floor and
        cap) or before (the irrigation before the day's ET and rain). kcf scales the crop coefficients."""
        D = np.zeros(self.n) if D0 is None else D0.copy()
        ac = self.acres if acres is None else acres
        eff = (self.ea if ea is None else ea) if use_ea else 1.0
        upto = upto or K.WEATHER_END
        for d in K.daterange(start, upto):
            e, r = self.wx[d]
            net = water[d] * 12.0 / ac * eff if d in water else 0.0
            et = kcf * self.kc[:, d.month - 1] * e
            if order == "net":
                D = np.maximum(D + et - r - net, 0.0)
            elif order == "after":
                D = np.maximum(D + et - r, 0.0)
                D = np.minimum(D, self.taw) if cap else D
                D = np.maximum(D - net, 0.0)
            else:
                D = np.maximum(D - net, 0.0)
                D = np.maximum(D + et - r, 0.0)
            D = np.minimum(D, self.taw) if cap else D
        return D

    # -------------------------------------------------------------------------------------------- the stations
    def station_path(self, st, ea_value, water=None, kcf=1.0):
        """The record run for one station's field at one efficiency: {date: depletion at the end of the day} from the
        end of 28 January 2026 (field capacity at the end of 27 January) to the end of 31 January 2027."""
        water = water if water is not None else self.water26()
        i = self.index[st["field"]]
        kc, taw, ac = self.kc[i], self.taw[i], self.acres[i]
        D, out = 0.0, {}
        for d in K.daterange(K.WINTER_FULL + dt.timedelta(days=1), K.WEATHER_END):
            e, r = self.wx[d]
            w = water[d][i] * 12.0 / ac * ea_value if d in water else 0.0
            D = min(taw, max(0.0, D + kcf * kc[d.month - 1] * e - r - w))
            out[d] = D
        return out

    def station_sse(self, st, ea_value, days=None, shift=0, kcf=1.0):
        """Squared misfit of the record run at ea_value against the station's readings (full point less profile water,
        each reading the end of the day before its local midnight); shift compares a reading with the run that many
        days later (a reading's UTC date taken as its day: shift 1)."""
        path = self.station_path(st, ea_value, kcf=kcf)
        tot = 0.0
        for d, v in st["reads"].items():
            dd = d + dt.timedelta(days=shift)
            if dd not in path or (days and not (days[0] <= d <= days[1])):
                continue
            tot += (st["full"] - v - path[dd]) ** 2
        return tot

    def station_sse_grid(self, st, grid, days=None, shift=0, kcf=1.0):
        """station_sse at every efficiency in grid at once (the run vectorised over the efficiencies)."""
        water = self.water26()
        i = self.index[st["field"]]
        kc, taw, ac = self.kc[i], self.taw[i], self.acres[i]
        D = np.zeros(len(grid))
        path = {}
        for d in K.daterange(K.WINTER_FULL + dt.timedelta(days=1), K.WEATHER_END):
            e, r = self.wx[d]
            w = water[d][i] * 12.0 / ac * grid if d in water else 0.0
            D = np.minimum(taw, np.maximum(0.0, D + kcf * kc[d.month - 1] * e - r - w))
            path[d] = D
        tot = np.zeros(len(grid))
        for d, v in st["reads"].items():
            dd = d + dt.timedelta(days=shift)
            if dd not in path or (days and not (days[0] <= d <= days[1])):
                continue
            tot += (st["full"] - v - path[dd]) ** 2
        return tot

    def calibrate_station(self, st, lo=0.55, hi=1.0, **kw):
        """The efficiency that brings the record run closest to the station's readings (least squares): the best of a
        grid at a ten-thousandth, then golden section to a millionth between its neighbours."""
        grid = np.round(np.arange(lo, hi + 1e-12, 0.0001), 4)
        k = int(np.argmin(self.station_sse_grid(st, grid, **kw)))
        a, b = grid[max(k - 1, 0)], grid[min(k + 1, len(grid) - 1)]
        g = (5 ** 0.5 - 1) / 2
        c, d_ = b - g * (b - a), a + g * (b - a)
        fc, fd = self.station_sse(st, c, **kw), self.station_sse(st, d_, **kw)
        while b - a > 1e-6:
            if fc <= fd:
                b, d_, fd = d_, c, fc
                c = b - g * (b - a)
                fc = self.station_sse(st, c, **kw)
            else:
                a, c, fc = c, d_, fd
                d_ = a + g * (b - a)
                fd = self.station_sse(st, d_, **kw)
        return float((a + b) / 2)

    def station_efficiency(self, how="mean", **kw):
        """Each station's efficiency and the efficiency by method (drip, micro-sprinkler and sprinkler stations; a
        border field's efficiency stays Table 4's, its sets running past field capacity)."""
        per = {st["station"]: self.calibrate_station(st, **kw) for st in self.stations if st["method"] != "Border"}
        by = {}
        for m in ("Drip", "Micro-sprinkler", "Sprinkler"):
            v = [per[st["station"]] for st in self.stations if st["method"] == m]
            by[m] = float(np.median(v)) if how == "median" else float(np.mean(v))
        return per, by

    def ea_from(self, by, border=None):
        return np.array([by.get(f["method"], self.T["ea"][f["method"]]) if f["method"] != "Border" else
                         (border if border is not None else self.T["ea"]["Border"]) for f in self.fields])

    def water26(self):
        if not hasattr(self, "_water26"):
            self._water26 = self.set_water()
        return self._water26

    def start_golden(self):
        """The golden start: each field's root zone as it stands at the end of 31 January 2027, its 2026 season run
        through the balance from field capacity at the end of 27 January 2026 (the 2025-26 winter's rains had every
        root zone full) on the water its delivered sets put in, at the share of that water that reaches the root zone
        as the soil moisture stations' 2026 readings show it for the field's irrigation method."""
        if not hasattr(self, "ea_by"):
            self.ea_station, self.ea_by = self.station_efficiency()
        self.ea_record = self.ea_from(self.ea_by)
        D = self.hindcast(self.water26(), ea=self.ea_record)
        self.D_golden = D
        return lambda y: D

    def start_t4(self):
        """The default path: the 2026 record run at the Table 4 efficiency (WOP 3.5's efficiency read across to the
        record)."""
        D = self.hindcast(self.water26())
        self.D_t4 = D
        return lambda y: D

    def station_jan(self):
        """{field index: the station's depletion at the end of 31 January 2027 (full point less profile water)}."""
        out = {}
        for st in self.stations:
            out[self.index[st["field"]]] = st["full"] - st["reads"][K.WEATHER_END]
        return out

    def start_p1(self):
        """The stations' 31 January readings on their own fields, the Table 4 run on every other field."""
        D = self.hindcast(self.water26()).copy()
        for i, v in self.station_jan().items():
            D[i] = v
        return lambda y: D

    def start_offset(self):
        """The Table 4 run shifted, method by method, by the stations' mean 31 January offset from it."""
        D4 = self.hindcast(self.water26())
        jan = self.station_jan()
        D = D4.copy()
        meth = np.array([f["method"] for f in self.fields])
        for m in K.METHODS:
            ix = [i for i in jan if meth[i] == m]
            off = float(np.mean([jan[i] - D4[i] for i in ix]))
            D[meth == m] = np.clip(D4[meth == m] + off, 0.0, self.taw[meth == m])
        for i, v in jan.items():
            D[i] = v
        return lambda y: D

    def start_station_share(self):
        """The stations' 31 January depletion as a share of total available water, method by method, on every field."""
        jan = self.station_jan()
        meth = np.array([f["method"] for f in self.fields])
        D = np.zeros(self.n)
        for m in K.METHODS:
            ix = [i for i in jan if meth[i] == m]
            share = float(np.mean([jan[i] / self.taw[i] for i in ix]))
            D[meth == m] = np.minimum(share, 1.0) * self.taw[meth == m]
        for i, v in jan.items():
            D[i] = v
        return lambda y: D

    def start_held_taw(self):
        """Every drip, micro-sprinkler and sprinkler field at its total available water, every border field from its
        last set taken as a refill."""
        D3 = self.carry([self.last_delivery[i] for i in self.ids])
        meth = np.array([f["method"] for f in self.fields])
        D = np.where(meth != "Border", self.taw, D3)
        return lambda y: D

    def start_partial(self):
        """The stations' efficiency on the stations' crops (almonds, wine grapes, pasture, alfalfa), Table 4 on the
        others."""
        self.start_golden()
        crops = {st["crop"] for st in self.stations}
        ea = np.array([self.ea_record[i] if f["crop"] in crops else self.ea[i] for i, f in enumerate(self.fields)])
        D = self.hindcast(self.water26(), ea=ea)
        return lambda y: D

    def start_regression(self):
        """A fitted model on the obvious covariates: the stations' 31 January depletion as a share of total available
        water regressed (least squares) on the day of the year of each field's last 2026 set, its crop coefficient in
        January and whether it is a drip, micro-sprinkler or sprinkler field, carried to every field."""
        jan = self.station_jan()
        ix = list(jan)
        last = np.array([self.last_delivery[i].timetuple().tm_yday for i in self.ids], dtype=float)
        held = np.array([1.0 if f["method"] != "Border" else 0.0 for f in self.fields])
        X = np.array([np.ones(self.n), last, self.kc[:, 0], held]).T
        y = np.array([jan[i] / self.taw[i] for i in ix])
        beta = np.linalg.lstsq(X[ix], y, rcond=None)[0]
        D = np.clip(X @ beta, 0.0, 1.0) * self.taw
        for i, v in jan.items():
            D[i] = v
        return lambda y: D

    def kc_factor(self):
        """The stations' 2026 readings fitted by a factor on the crop coefficients at the Table 4 efficiency."""
        best = None
        for k in range(800, 1101):
            kf = k / 1000.0
            sse = sum(self.station_sse(st, self.T["ea"][st["method"]], kcf=kf) for st in self.stations
                      if st["method"] != "Border")
            if best is None or sse < best[0]:
                best = (sse, kf)
        return best[1], best[0]

    def start_kc_fit(self):
        kf, _ = self.kc_factor()
        D = self.hindcast(self.water26(), kcf=kf)
        return lambda y: D

    def start_first_set(self):
        """R3c: the 2026 season run on the record from each field's first delivered set, taken as returning it to field
        capacity (at the stations' efficiency)."""
        self.start_golden()
        water = self.water26()
        first = {}
        for fid, a, b, cfs, status in self.record:
            if status == "Delivered":
                first[fid] = min(first.get(fid, b.date()), b.date())
        a_ord = np.array([first[i].toordinal() for i in self.ids])
        D = np.zeros(self.n)
        for d in K.daterange(K.WINTER_FULL + dt.timedelta(days=1), K.WEATHER_END):
            e, r = self.wx[d]
            net = water[d] * 12.0 / self.acres * self.ea_record if d in water else 0.0
            nd = np.minimum(np.maximum(D + self.kc[:, d.month - 1] * e - r - net, 0.0), self.taw)
            D = np.where(d.toordinal() > a_ord, nd, np.where(d.toordinal() == a_ord, 0.0, nd))
        return lambda y: D

    def start_gross(self):
        D = self.hindcast(self.water26(), use_ea=False)
        return lambda y: D

    def start_cancelled(self):
        self.start_golden()
        D = self.hindcast(self.set_water(statuses=("Delivered", "Cancelled")), ea=self.ea_record)
        return lambda y: D

    def start_canal_close(self):
        D = self.carry([dt.date(2026, 10, 31)] * self.n)
        return lambda y: D

    def start_nocap(self):
        self.start_golden()
        D = self.hindcast(self.water26(), cap=False, ea=self.ea_record)
        return lambda y: D

    def _model_season_then_winter(self, year, D):
        for d in K.season_days(year):
            D = self.day(D, d)
            D = np.where(D >= self.raw, 0.0, D)
        for d in K.daterange(dt.date(year, 11, 1), dt.date(year + 1, 1, 31)):
            D = self.day(D, d)
        return D

    def start_own_winter(self):
        """R1: each weather year from the state its own preceding autumn and winter left, the replay rule run through
        the year before from a full root zone on 1 February."""
        cache = {}

        def f(y):
            if y not in cache:
                cache[y] = self._model_season_then_winter(y - 1, np.zeros(self.n))
            return cache[y]
        return f

    def start_model_2026(self):
        """R2: the replay rule run through the 2026 season (from the state the 2025-26 winter left, itself from the
        rule run through 2025), then the observed autumn and winter."""
        D25 = self._model_season_then_winter(2025, np.zeros(self.n))
        D = self._model_season_then_winter(2026, D25)
        return lambda y: D


def order_from(req):
    """WOP 2.3: the smallest whole number of thousands covering the requirement in at least 24 of 30 years."""
    ranked = sorted(req.items(), key=lambda kv: kv[1])
    y24, v24 = ranked[23]
    return dict(order=K.ceil_to(v24, K.ORDER_STEP), year=y24, req=v24, ranked=ranked)


def schedule(order, mon):
    """WOP 5.1: shares of the thirty years' summed monthly requirements, each month to the nearest 100 acre-feet,
    October the difference."""
    tot_m = np.zeros(9)
    for y in K.REPLAY_YEARS:
        tot_m += mon[y]
    shares = tot_m / tot_m.sum()
    raw = order * shares
    months = [K.round_to(x, K.MONTH_STEP) for x in raw[:8]]
    months.append(order - sum(months))
    return months, raw, shares


def headgate_from_farm_by_lateral(farm_by_lat, loss, main):
    return float((np.asarray(farm_by_lat) / (1 - np.asarray(loss))).sum() / (1 - main))


# ---------------------------------------------------------------------------------------------------- loading the pack

def read_csv(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.reader(fh))


def load_tables(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    T = dict(kc={}, root={}, mad={}, soils={}, ea={}, loss={}, main_loss=None)

    def rows(ws):
        out = []
        for row in ws.iter_rows(values_only=True):
            out.append(row)
        return out

    for row in rows(wb["T1 Crop coefficients"]):
        if row and isinstance(row[0], str) and row[0] in K.CROPS and isinstance(row[2], (int, float)):
            T["kc"][row[0]] = [float(x) for x in row[2:14]]
    for row in rows(wb["T2 Crops"]):
        if row and isinstance(row[0], str) and row[0] in K.CROPS and isinstance(row[2], (int, float)):
            T["root"][row[0]] = float(row[2])
            T["mad"][row[0]] = float(row[3])
    for row in rows(wb["T3 Soil units"]):
        if row and row[0] is not None and str(row[0]) in K.SOILS and isinstance(row[2], (int, float)):
            lim = row[3]
            T["soils"][str(row[0])] = (float(row[2]), None if lim in (None, "", "none", "None") else float(lim))
    for row in rows(wb["T4 Application efficiency"]):
        if row and row[0] in K.METHODS and isinstance(row[1], (int, float)):
            T["ea"][row[0]] = float(row[1])
    for row in rows(wb["T5 Loss allowances"]):
        if row and isinstance(row[0], str) and row[0].startswith("Lateral ") and isinstance(row[1], (int, float)):
            T["loss"]["L" + row[0].split()[1]] = float(row[1]) / 100.0
        if row and isinstance(row[0], str) and row[0].startswith("Main canal") and isinstance(row[1], (int, float)):
            T["main_loss"] = float(row[1]) / 100.0
    return T


def load_weather(target):
    s39 = {}
    rows = read_csv(os.path.join(target, K.ST39))
    head = rows[0]
    i_date, i_eto, i_pr = head.index("Date"), head.index("ETo (mm)"), head.index("Precip (mm)")
    for r in rows[1:]:
        m, d_, y = (int(x) for x in r[i_date].split("/"))
        s39[dt.date(y, m, d_)] = (float(r[i_eto]) / 25.4, float(r[i_pr]) / 25.4)
    yard = {}
    rows = read_csv(os.path.join(target, K.YARD))
    head = rows[0]
    i_ts, i_e, i_r = head.index("Date"), head.index("ETo_in"), head.index("Rain_in")
    for r in rows[1:]:
        if r[i_e] == "" or r[i_r] == "":
            continue
        d = dt.date.fromisoformat(r[i_ts])
        yard[d] = (float(r[i_e]), float(r[i_r]))
    wx = {}
    for d in K.daterange(K.WEATHER_START, K.WEATHER_END):
        wx[d] = yard[d] if (d >= K.SPLICE and d in yard) else s39[d]
    return wx, s39, yard


def load_requests(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Requests"]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "Request no.")
    head = list(rows[hi])
    out = []
    for r in rows[hi + 1:]:
        if not r or r[0] is None:
            continue
        out.append(dict(zip(head, r)))
    return out


def load(target):
    T = load_tables(os.path.join(target, K.WOP_TABLES))
    wx, s39, yard = load_weather(target)
    reqs = load_requests(os.path.join(target, K.REQUESTS))
    reg = {}
    rows = read_csv(os.path.join(target, K.REGISTER))
    head = rows[0]
    for r in rows[1:]:
        reg[r[head.index("field_id")]] = dict(zip(head, r))
    crop_by_name = {v: k for k, v in K.CROP_NAME.items()}
    fields = []
    for q in reqs:
        if q["Status"] != "Accepted":
            continue
        fid = q["Field"]
        g = reg[fid]
        crop = q["Crop code"]
        fields.append(dict(id=fid, lateral="L" + g["lateral"].split()[-1], soil=g["soil_unit"], crop=crop,
                           method=q["Irrigation method"], acres=float(q["Acres"]), turnout=g["turnout_id"]))
    fields.sort(key=lambda f: f["id"])
    # the orders: each field's last delivered 2026 set, each turnout's, and the latest cancelled order after it
    rows = read_csv(os.path.join(target, K.ORDERS))
    head = rows[0]
    ix = {c: head.index(c) for c in head}
    last, tlast, canc, t_of, record = {}, {}, {}, {}, []
    for r in rows[1:]:
        if r[ix["season"]] != "2026":
            continue
        a = dt.datetime.fromisoformat(r[ix["set_start"]])
        b = dt.datetime.fromisoformat(r[ix["set_end"]])
        day, end_day = a.date(), b.date()
        fid, tid = r[ix["field_id"]], r[ix["turnout"]]
        t_of[fid] = tid
        record.append((fid, a, b, float(r[ix["ordered_cfs"]]), r[ix["status"]]))
        if r[ix["status"]] == "Delivered":
            last[fid] = max(last.get(fid, end_day), end_day)
            tlast[tid] = max(tlast.get(tid, end_day), end_day)
        elif r[ix["status"]] == "Cancelled":
            canc[fid] = max(canc.get(fid, day), day)
    turnout_last = {f["id"]: tlast[t_of[f["id"]]] for f in fields}
    M = Model(fields, wx, T, last_delivery=last, turnout_last=turnout_last, cancelled_last=canc, record=record)
    M.s39, M.yard = s39, yard
    M.requests = reqs
    M.register = reg
    M.stations = load_stations(target, reg, {f["id"]: f for f in fields})
    return M


def load_stations(target, register, fields):
    """The soil moisture network's sites, keyed to the register's fields by APN, and each station's readings by the
    day each ends (a reading at local midnight ends the day before; the platform stamps UTC)."""
    import openpyxl
    from zoneinfo import ZoneInfo
    path = os.path.join(target, K.STATIONS)
    if not os.path.exists(path):
        return []
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = list(wb["Sites"].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "Station")
    head = list(rows[hi])
    by_apn = {g["apn"].replace("-", ""): fid for fid, g in register.items()}
    crop_code = {v: k for k, v in K.CROP_NAME.items()}
    sites = {}
    for r in rows[hi + 1:]:
        if not r or r[0] is None:
            continue
        x = dict(zip(head, r))
        fid = by_apn[str(x["APN"]).strip()]
        sites[x["Station"]] = dict(station=x["Station"], field=fid, method=x["Irrigation"], crop=crop_code[x["Crop"]],
                                   full=float(x["Full point (in)"]), depth=float(x["Profile to (in)"]), reads={})
    local = ZoneInfo("America/Los_Angeles")
    for r in wb["Daily"].iter_rows(min_row=2, values_only=True):
        if not r or r[0] is None:
            continue
        t = dt.datetime.strptime(r[1], "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=dt.timezone.utc).astimezone(local)
        sites[r[0]]["reads"][t.date() - dt.timedelta(days=1)] = float(r[2])
    return [sites[k] for k in sorted(sites) if sites[k]["field"] in fields]


# ---------------------------------------------------------------------------------------------------- the ladder

RUNGS = [("R0", "full on 1 February", "start_full"),
         ("R1", "each weather year's own preceding autumn and winter", "start_own_winter"),
         ("R2", "the replay rule run through 2026, then the observed winter", "start_model_2026"),
         ("R2b", "full at the canal's close (31 October 2026), then the observed winter", "start_canal_close"),
         ("R3", "each field's last delivered 2026 set taken as a refill, then the observed winter", "start_last_set"),
         ("R4g", "the 2026 record run with all of each set's delivered water reaching the root zone", "start_gross"),
         ("KC", "the stations' 2026 readings fitted by a crop coefficient factor at the Table 4 efficiency",
          "start_kc_fit"),
         ("R3c", "the record run at the stations' efficiency from each field's first set taken as a refill",
          "start_first_set"),
         ("R2d", "the record run at the stations' efficiency with cancelled sets counted as water", "start_cancelled"),
         ("G", "the 2026 record run at the efficiency the stations' 2026 readings show (golden)", "start_golden"),
         ("PT", "the stations' efficiency on the stations' crops, Table 4 on the rest", "start_partial"),
         ("AB", "the Table 4 run shifted by the stations' 31 January offset, method by method", "start_offset"),
         ("P1", "the stations' 31 January readings on their own fields, the Table 4 run elsewhere", "start_p1"),
         ("T4", "the 2026 record run at the Table 4 efficiency, checked against the stations' January readings "
                "(the default path)", "start_t4"),
         ("MI", "the stations' 31 January share of total available water on every field of the method",
          "start_station_share"),
         ("RG", "the stations' 31 January share regressed on the last set's day, the January crop coefficient and the "
                "method", "start_regression"),
         ("HT", "every drip, micro-sprinkler and sprinkler field at its total available water", "start_held_taw"),
         ("O2", "the golden start with the root-zone cap dropped", "start_nocap")]


def ladder(M):
    out = {}
    for key, label, meth in RUNGS:
        req, mon = M.thirty(getattr(M, meth)())
        o = order_from(req)
        months, raw, shares = schedule(o["order"], mon)
        out[key] = dict(label=label, reqs=req, mon=mon, months=months, raw_months=raw, **o)
    # O1: each full-start replay plus each field's 31 January depletion taken as a flat addition at the turnouts
    D = M.D_golden
    add_farm = D / M.ea * M.acres / 12.0
    add_lat = np.bincount(M.lat, weights=add_farm, minlength=9)
    add_head = headgate_from_farm_by_lateral(add_lat, M.loss, M.main)
    req = {y: v + add_head for y, v in out["R0"]["reqs"].items()}
    o = order_from(req)
    out["O1"] = dict(label="the full-start replays plus each field's 31 January depletion as a flat addition",
                     reqs=req, mon=None, months=None, raw_months=None, **o)
    # O3: the stations' efficiency carried into the replays' deliveries as well (WOP 3.5 files Table 4's)
    M3 = Model(M.fields, M.wx, M.T, last_delivery=M.last_delivery, record=M.record)
    M3.ea = M.ea_record.copy()
    req = M3.thirty(lambda y: D)[0]
    o = order_from(req)
    out["O3"] = dict(label="the golden start with the stations' efficiency in the replays' deliveries too", reqs=req,
                     mon=None, months=None, raw_months=None, **o)
    return out


# ---------------------------------------------------------------------------------------------------- the asks

def tid_from_book(x):
    a, b = x.split("-")
    return "%02d-%03d" % (int(a), int(b))


def tid_from_gis(x):
    return "%s-%s" % (x[1:3], x[3:6])


def load_asks(target):
    import openpyxl
    A = {}
    wb = openpyxl.load_workbook(os.path.join(target, K.SHOP_LOG), data_only=True)
    units, cap = {}, {}
    rows = list(wb["Meter inventory"].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "Meter")
    for r in rows[hi + 1:]:
        if not r or r[0] is None:
            continue
        units[r[0]] = "AF" if r[3] == "acre-feet" else "HCF"
        cap[r[0]] = float(str(r[4]).replace(",", ""))
    changes = []
    rows = list(wb["Work orders 2026"].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "WO no.")
    head = list(rows[hi])
    for r in rows[hi + 1:]:
        if not r or r[0] is None:
            continue
        x = dict(zip(head, r))
        if not x["Meter out"]:
            continue
        day = x["Date"].date() if isinstance(x["Date"], dt.datetime) else x["Date"]
        hh, mm = (int(v) for v in str(x["Time"]).split(":"))
        changes.append(dict(turnout=tid_from_gis(x["Turnout"]), when=dt.datetime(day.year, day.month, day.day, hh, mm),
                            old=x["Meter out"], removal=float(x["Reg. out"]), new=x["Meter in"],
                            installation=float(x["Reg. in"]), unit=x["Unit"]))
    A["units"], A["cap"], A["changes"] = units, cap, changes
    reads = {}
    rows = read_csv(os.path.join(target, K.METER_BOOK))
    head = rows[0]
    ix = {c: head.index(c) for c in head}
    for r in rows[1:]:
        m, d_, y = (int(x) for x in r[ix["read_date"]].split("/"))
        hh, mm = (int(x) for x in r[ix["read_time"]].split(":"))
        tid = tid_from_book(r[ix["turnout"]])
        reads.setdefault(tid, []).append((dt.datetime(y, m, d_, hh, mm), r[ix["meter"]], float(r[ix["reading"]]),
                                          r[ix["reg_units"]]))
    for v in reads.values():
        v.sort()
    A["reads"] = reads
    lat = {}
    rows = read_csv(os.path.join(target, K.REGISTER))
    head = rows[0]
    for r in rows[1:]:
        lat[tid_from_gis(r[head.index("turnout_id")])] = "L" + r[head.index("lateral")].split()[-1]
    A["lateral"] = lat
    heads = []
    rows = read_csv(os.path.join(target, K.HEADINGS))
    head = rows[0]
    for r in rows[1:]:
        heads.append((dt.date.fromisoformat(r[head.index("date")]), "L" + r[head.index("lateral")].split()[-1],
                      float(r[head.index("volume_af")])))
    A["headings"] = heads
    A["estimates"] = heading_estimates(target)
    wb = openpyxl.load_workbook(os.path.join(target, K.STRUCTURES), data_only=True)
    weir = {}
    rows = list(wb["Tail weirs"].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "Lateral")
    for r in rows[hi + 1:]:
        if not r or r[0] is None:
            continue
        weir["L" + r[0].split()[-1]] = (float(r[2]), float(r[3]), float(r[4]))
    A["weir"] = weir
    sheets = {}
    rows = read_csv(os.path.join(target, K.WEIRS))
    for r in rows[1:]:
        m, d_, y = (int(x) for x in r[0].split("/"))
        sheets.setdefault("L" + r[1].split()[-1], []).append((dt.date(y, m, d_), float(r[4]), r[5], r[3]))
    for v in sheets.values():
        v.sort()
    A["sheets"] = sheets
    return A


def heading_estimates(target):
    """The Water Master's heading estimates for the days a heading logger kept no total, from the operations log:
    {(lateral, date): AF}."""
    import docx
    import re
    t = docx.Document(os.path.join(target, K.OPS_LOG)).tables[0]
    months = {m.lower(): k for k, m in enumerate(["January", "February", "March", "April", "May", "June", "July",
                                                  "August", "September", "October", "November", "December"], 1)}
    abbr = {m[:3]: k for m, k in months.items()}
    mon = r"(January|February|March|April|May|June|July|August|September|October|November|December|Jan|Feb|Mar|Apr|" \
          r"Jun|Jul|Aug|Sep|Oct|Nov|Dec)"
    out = {}
    for row in t.rows[1:]:
        c = [x.text for x in row.cells]
        day = dt.date.fromisoformat(c[0])
        text = c[2]
        if "estimat" not in text.lower():
            continue
        lat = re.search(r"\b(?:Lateral|Lat|L)\s?(\d)\b", text)
        pairs = []
        for a, b, v in re.findall(r"\b(\d{1,2}) %s (\d+\.\d)" % mon, text):
            pairs.append((months.get(b.lower(), abbr.get(b.lower()[:3])), int(a), float(v)))
        for b, a, v in re.findall(r"\b%s (\d{1,2}) (\d+\.\d)" % mon, text):
            pairs.append((months.get(b.lower(), abbr.get(b.lower()[:3])), int(a), float(v)))
        for m_, a, v in re.findall(r"\b(\d{1,2})/(\d{1,2}) (\d+\.\d)", text):
            pairs.append((int(m_), int(a), float(v)))
        for m_, a, v in pairs:
            out[("L" + lat.group(1), dt.date(day.year, m_, a))] = v
    return out


MOVED_FROM = dt.date(2026, 6, 15)


def turnout_intervals(A, tid, mode="chain"):
    """Acre-feet a turnout's meters recorded between consecutive reads, as (end of interval, AF).
    mode chain: within a meter's own reads, and across a change through the shop's removal and installation reads;
    mode serial: within a meter's own reads only; mode turnout: consecutive reads whatever the meter."""
    rd = A["reads"][tid]
    ch = [c for c in A["changes"] if c["turnout"] == tid]
    out = []

    def af(serial, a, b):
        c = A["cap"][serial]
        d = (b - a) % c
        return d / 435.6 if A["units"][serial] == "HCF" else d

    for (t0, s0, v0, _u0), (t1, s1, v1, _u1) in zip(rd, rd[1:]):
        if s0 == s1:
            out.append((t1, af(s0, v0, v1)))
        elif mode == "chain":
            c = next(c for c in ch if c["old"] == s0 and c["new"] == s1)
            out.append((t1, af(s0, v0, c["removal"]) + af(s1, c["installation"], v1)))
        elif mode == "turnout":
            out.append((t1, (v1 - v0) % A["cap"][s1] if A["units"][s1] == "AF" else ((v1 - v0) % A["cap"][s1]) / 435.6))
        # mode serial: the interval across the change is lost
    return out


def ask_deliveries(A, mode="chain", as_of=True):
    """Each lateral's 2026 deliveries: every turnout's metered water on each read interval, on the lateral that fed
    the turnout over that interval (the west spur on Lateral 5 through its 15 June read, Lateral 6 after)."""
    deliv = {l: 0.0 for l in K.LATERALS}
    for tid in A["reads"]:
        reg_lat = A["lateral"][tid]
        built = "L%d" % int(tid[:2])
        for t1, v in turnout_intervals(A, tid, mode):
            if as_of and built != reg_lat:
                lat = built if t1.date() <= MOVED_FROM else reg_lat
            else:
                lat = reg_lat
            deliv[lat] += v
    return deliv


def ask_inflow(A, estimates=True):
    """Each lateral's 2026 water put in at its heading: the logger's daily totals and, for the days a logger kept no
    total, the Water Master's estimates."""
    inflow = {l: 0.0 for l in K.LATERALS}
    for d, l, v in A["headings"]:
        inflow[l] += v
    if estimates:
        for (l, d), v in A["estimates"].items():
            inflow[l] += v
    return inflow


def ask_a(A, mode="chain", as_of=True, estimates=True):
    deliv = ask_deliveries(A, mode, as_of)
    inflow = ask_inflow(A, estimates)
    share = {l: 100.0 * (inflow[l] - deliv[l]) / inflow[l] for l in K.LATERALS}
    return deliv, share, inflow


def board_counts(A, lat):
    """Boards in place at each reading, from the riders' remarks (a board set, added or pulled before that day's
    reading, its time noted)."""
    import re
    n = 0
    out = {}
    for d, g, rem, rt in A["sheets"][lat]:
        rem = (rem or "").lower().strip()
        if re.match(r"^set 2in brd \d{4}$", rem) or re.match(r"^added 2in brd \d{4}$", rem):
            n += 1
        elif re.match(r"^pulled 1 brd \d{4}$", rem):
            n -= 1
        elif re.match(r"^brds out \d{4}$", rem):
            n = 0
        out[d] = n
    return out


def ask_b(A, boards=True, fill=True):
    spill = {}
    for lat in K.LATERALS:
        L, crest, C_ = A["weir"][lat]
        bc = board_counts(A, lat)
        rd = {d: g for d, g, _, _ in A["sheets"][lat]}
        tot = 0.0
        cur = None
        for d in K.season_days(2026):
            if d in rd:
                cur = (rd[d], bc[d])
            elif not fill:
                continue
            g, nb = cur
            h = g - (crest + (2.0 / 12.0 * nb if boards else 0.0))
            if h > 0:
                tot += C_ * L * h ** 1.5 * K.AF_PER_CFS_DAY
        spill[lat] = tot
    return spill
