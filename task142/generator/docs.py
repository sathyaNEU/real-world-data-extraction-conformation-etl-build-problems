"""The pack's documents and workbooks: Schedule 2 of the Bellbird agreement (the January to March 2027 block), the
program streams and visit schedule, the bone and joint infection pathway, the members' 2026-27 operating plans
(revised October 2026), the in-house capacity return, Schedules 3 and 4 (fees and service levels), Bellbird's intake
register, the data dictionary, and the two documents nothing in the solution reads (the 2025-26 patient experience
survey and Bellbird's 2025-26 service report)."""
from __future__ import annotations

import datetime as dt

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import A4                                                # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import cm                                                    # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import common as K   # noqa: E402
import world as WD   # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/"
for _n, _f in (("Carlito", "crosextra/Carlito-Regular.ttf"), ("Carlito-Bold", "crosextra/Carlito-Bold.ttf"),
               ("Carlito-Italic", "crosextra/Carlito-Italic.ttf"),
               ("LiberationSerif", "liberation/LiberationSerif-Regular.ttf"),
               ("LiberationSerif-Bold", "liberation/LiberationSerif-Bold.ttf")):
    pdfmetrics.registerFont(TTFont(_n, FONT_DIR + _f))
SANS, SANS_B, SANS_I = "Carlito", "Carlito-Bold", "Carlito-Italic"
SERIF, SERIF_B = "LiberationSerif", "LiberationSerif-Bold"
INK = colors.HexColor("#1f2a30")
GREY = colors.HexColor("#5b646b")
GREEN = colors.HexColor("#2f5d50")
RULE = colors.HexColor("#a9b4ad")
PALE = colors.HexColor("#e7eeea")

BODY = ParagraphStyle("body", fontName=SANS, fontSize=10.2, leading=13.4, textColor=INK, spaceAfter=5)
CLAUSE = ParagraphStyle("clause", parent=BODY, leftIndent=26, firstLineIndent=-26)
H1 = ParagraphStyle("h1", fontName=SANS_B, fontSize=15, leading=18.5, textColor=GREEN, spaceAfter=3)
H1S = ParagraphStyle("h1s", fontName=SANS, fontSize=9.8, leading=12.8, textColor=GREY, spaceAfter=9)
H2 = ParagraphStyle("h2", fontName=SANS_B, fontSize=11.2, leading=14, textColor=INK, spaceBefore=8, spaceAfter=3)
NOTE = ParagraphStyle("note", fontName=SANS_I, fontSize=8.6, leading=11, textColor=GREY, spaceBefore=4)
CELL = ParagraphStyle("cell", fontName=SANS, fontSize=8.9, leading=10.8, textColor=INK)
CELLB = ParagraphStyle("cellb", fontName=SANS_B, fontSize=8.9, leading=10.8, textColor=INK)
LBODY = ParagraphStyle("lbody", fontName=SERIF, fontSize=10.4, leading=13.8, textColor=INK, spaceAfter=5)
LCLAUSE = ParagraphStyle("lclause", parent=LBODY, leftIndent=28, firstLineIndent=-28)
LH1 = ParagraphStyle("lh1", fontName=SERIF_B, fontSize=13.5, leading=17, textColor=INK, spaceAfter=4)
LH2 = ParagraphStyle("lh2", fontName=SERIF_B, fontSize=11, leading=14, textColor=INK, spaceBefore=8, spaceAfter=3)
LSUB = ParagraphStyle("lsub", fontName=SERIF, fontSize=9.6, leading=12.4, textColor=GREY, spaceAfter=8)


def money(x):
    return "${:,.2f}".format(x)


def _footer(left, right, font=SANS):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 7.8)
        canvas.setFillColor(GREY)
        canvas.drawString(2.0 * cm, 1.2 * cm, left)
        canvas.drawRightString(A4[0] - 2.0 * cm, 1.2 * cm, "%s  |  page %d" % (right, doc.page))
        canvas.restoreState()
    return draw


def _doc(path, title, author):
    return SimpleDocTemplate(path, pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=1.9 * cm,
                             bottomMargin=2.0 * cm, title=title, author=author, subject="", creator="", producer="")


def _table(rows, widths, head=True, style=None, font=None):
    cs, cb = (CELL, CELLB) if font is None else font
    data = [[Paragraph(str(c), cb if (head and i == 0) else cs) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, hAlign="LEFT")
    st = [("GRID", (0, 0), (-1, -1), 0.4, RULE), ("VALIGN", (0, 0), (-1, -1), "TOP"),
          ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]
    if head:
        st.append(("BACKGROUND", (0, 0), (-1, 0), PALE))
    t.setStyle(TableStyle(st + (style or [])))
    return t


# ---------------------------------------------------------------------------------------------------- Schedule 2

def write_schedule2(path):
    P = K.PEOPLE
    s = [Paragraph("Overflow Nursing Services Agreement", LSUB),
         Paragraph("Ironbark Health Alliance Inc. and Bellbird Home Nursing Pty Ltd", LSUB),
         Paragraph("Schedule 2: Overflow visit block for January to March 2027", LH1),
         Paragraph("Variation 9, executed 2 October 2026. This Schedule sets the block for the quarter from 1 January "
                   "2027 to 31 March 2027 and replaces Schedule 2 as varied on 3 July 2026 for that quarter. Terms with "
                   "capitals have the meanings in clause 1 of the Agreement.", LSUB)]

    def cl(no, text):
        s.append(Paragraph("%s&nbsp;&nbsp;&nbsp;%s" % (no, text), LCLAUSE))

    s.append(Paragraph("The block", LH2))
    cl("2.1", "Bellbird will hold 850 nurse visits in reserve for the Members' Hospital in the Home patients for the "
              "quarter from 1 January 2027 to 31 March 2027 (the Quarter).")
    cl("2.2", "For 2026-27 the Members book the block in this order: Northgate Health, Riverbend Health, Westmere "
              "Health, Kingsdale Health, Hillcrest Health.")
    s.append(Paragraph("Each Member's lots", LH2))
    cl("2.3", "A Member's requirement for the Quarter is the number of nurse visits its admitted Hospital in the Home "
              "patients receive on the Quarter's days, whoever makes them, with patients admitted before the Quarter "
              "included for their days inside it.")
    cl("2.4", "Each Member books against the 90th percentile of its requirement for the Quarter, the visits of the "
              "Quarter taken together.")
    cl("2.5", "The expected value of a Member's requirement is the Member's Hospital in the Home nurse visits for "
              "January, February and March 2027 in its 2026-27 operating plan as revised in October 2026.")
    cl("2.6", "Each Member books the lots of 25 visits that cover the amount by which the 90th percentile of its "
              "requirement exceeds its in-house capacity for the Quarter in the Alliance's in-house capacity return. "
              "A Member whose 90th percentile does not exceed its in-house capacity books no lots.")
    cl("2.7", "Members book in the order in clause 2.2 until the block is used. A Member that needs more lots than "
              "remain takes the lots that remain, and the Members after it book none.")
    s.append(Paragraph("Lodging and use", LH2))
    cl("2.8", "The Alliance lodges the booking with Bellbird by 5.00 pm on Friday 6 November 2026, stating each "
              "Member's lots in visits.")
    cl("2.9", "Booked visits are held for the Member that booked them. Visits Bellbird makes in the Quarter are "
              "invoiced under Schedule 3, whether or not they fall within a booking.")
    cl("2.10", "Bellbird may release booked visits that a Member has not called on by the fourteenth day before the "
               "end of the Quarter, after notice to the Alliance.")
    s.append(Spacer(1, 14))
    sig = [["For Ironbark Health Alliance Inc.", "For Bellbird Home Nursing Pty Ltd"],
           ["%s, Program Manager, Hospital in the Home" % P["requester"],
            "%s, Account Manager, Health Services" % P["bellbird"]],
           ["Signed 2 October 2026 under delegation of the Alliance board",
            "Signed 2 October 2026 under delegation of the managing director"]]
    s.append(_table(sig, [8.2 * cm, 8.2 * cm], head=False, font=(ParagraphStyle("lc", fontName=SERIF, fontSize=9.4,
                                                                                leading=12, textColor=INK),
                                                                 ParagraphStyle("lcb", fontName=SERIF_B, fontSize=9.4,
                                                                                leading=12, textColor=INK))))
    foot = _footer("Overflow Nursing Services Agreement, Schedule 2 (Variation 9)", "IHA-BHN-2024-01", font=SERIF)
    doc = _doc(path, "Schedule 2: Overflow visit block for January to March 2027", P["requester"])
    doc.build(s, onFirstPage=foot, onLaterPages=foot)


# ---------------------------------------------------------------------------------------------------- Schedules 3 and 4

def write_schedules34(path):
    P = K.PEOPLE
    s = [Paragraph("Overflow Nursing Services Agreement", LSUB),
         Paragraph("Ironbark Health Alliance Inc. and Bellbird Home Nursing Pty Ltd", LSUB),
         Paragraph("Schedule 3: Fees. Schedule 4: Referrals and service levels", LH1),
         Paragraph("As varied on 3 July 2026 (Variation 8: fees from 1 July 2026). Terms with capitals have the "
                   "meanings in clause 1 of the Agreement.", LSUB)]

    def cl(no, text):
        s.append(Paragraph("%s&nbsp;&nbsp;&nbsp;%s" % (no, text), LCLAUSE))

    s.append(Paragraph("Schedule 3: Fees", LH2))
    cl("3.1", "Bellbird invoices the Alliance monthly in arrears for the visits it made in the month to each Member's "
              "patients. The Alliance pays Bellbird and recharges each Member its share.")
    cl("3.2", "From 1 July 2026 the fee for a visit is set by the kind of visit. Fees are GST-free.")
    rows = [["Kind of visit", "Code", "Fee"]] + [[K.FEE_NAME[c], c, money(K.FEE[c])] for c in ("IVT", "WND", "CLN")]
    s.append(_table(rows, [7.0 * cm, 2.0 * cm, 3.0 * cm], font=(ParagraphStyle("c3", fontName=SERIF, fontSize=9.6,
                                                                              leading=12, textColor=INK),
                                                               ParagraphStyle("c3b", fontName=SERIF_B, fontSize=9.6,
                                                                              leading=12, textColor=INK))))
    s.append(Spacer(1, 6))
    cl("3.3", "A loading is added to the fee for a visit that starts at these times: Monday to Friday before 7.00 am "
              "or at or after 6.00 pm, 25 per cent; Saturday, 40 per cent; Sunday, 60 per cent; a Victorian public "
              "holiday, 100 per cent. Only one loading applies to a visit, the highest that applies.")
    cl("3.4", "Where Bellbird's nurse attends and the patient is not at home, or does not accept the visit, the visit is "
              "an attempted visit and its fee is %s. No loading applies to an attempted visit." % money(K.ATTEMPT_FEE))
    cl("3.5", "No fee is payable for a visit cancelled before Bellbird's nurse attends.")
    cl("3.6", "Times and dates in this Schedule are Melbourne time. A visit belongs to the day, and to the month, on "
              "which it starts.")
    cl("3.7", "Bellbird's visit records are the basis of each invoice. The Alliance may query an invoice within 60 "
              "days of receiving it.")
    s.append(Paragraph("Schedule 4: Referrals and service levels", LH2))
    cl("4.1", "A referral is a Member's request that Bellbird take over a patient's Hospital in the Home visits from "
              "a stated date.")
    cl("4.2", "Bellbird accepts or declines each referral within two hours of the Member making it.")
    cl("4.3", "Bellbird makes its first visit to the patient of an accepted referral within 24 hours of accepting it. "
              "The first visit is Bellbird's first attendance at the patient's home after it accepts the referral, "
              "whether the visit is completed or the patient is not at home.")
    cl("4.4", "Bellbird's nurses follow the Member's treatment orders and the Alliance Hospital in the Home visit "
              "schedule for the patient's program stream, and hand the patient back to the Member's own nurses when "
              "the Member asks.")
    cl("4.5", "The Alliance reviews Bellbird's referral response with Bellbird each quarter. Times in this Schedule are "
              "Melbourne time.")
    s.append(Spacer(1, 10))
    s.append(Paragraph("Contacts for this Schedule: %s (Alliance, Hospital in the Home program); %s (Alliance finance, "
                       "provider payments); %s (Bellbird, account management)." % (P["requester"], P["finance"],
                                                                                  P["bellbird"]), LSUB))
    foot = _footer("Overflow Nursing Services Agreement, Schedules 3 and 4 (Variation 8)", "IHA-BHN-2024-01",
                   font=SERIF)
    doc = _doc(path, "Schedules 3 and 4: Fees, referrals and service levels", P["requester"])
    doc.build(s, onFirstPage=foot, onLaterPages=foot)


# ---------------------------------------------------------------------------------------------------- streams

DX_TEXT = {"H01": "L02, L03, L08 (abscess, cellulitis, other local infection of skin)",
           "H02": "N10, N30, N39 (pyelonephritis, cystitis, urinary tract infection)",
           "H03": "J15, J18, J22, J44 (pneumonia, lower respiratory infection, COPD with infection)",
           "H04": "M00, M46, M86, T84 (septic arthritis, vertebral and other osteomyelitis, prosthetic joint infection)",
           "H05": "G06, I33, J86, K75 (intracranial abscess, endocarditis, empyema, liver abscess)",
           "H06": "T81, L89, L97 and others (wound complications, pressure and leg ulcers)",
           "H07": "I26, I80 (pulmonary embolism, deep vein thrombosis)",
           "H08": "I50 (heart failure)"}
SCHED_TEXT = {"H06": "One visit on the day of admission and every second day after it",
              "H08": "Two visits on each of the first two days at home, then one visit a day"}


def write_streamdoc(path):
    s = [Paragraph("%s  |  Hospital in the Home clinical governance group" % K.ALLIANCE, H1S),
         Paragraph("Hospital in the Home program streams and visit schedule", H1),
         Paragraph("Effective for admissions from 1 July 2025. Approved by the clinical governance group on 16 June "
                   "2025. Applies to the Hospital in the Home programs of the five Members.", H1S)]
    s.append(Paragraph("Program streams", H2))
    s.append(Paragraph("Every admission to a Member's Hospital in the Home program is recorded in one program stream, "
                       "set at admission from the principal diagnosis. The eight streams below replace the four "
                       "programs used for admissions to 30 June 2025 (AB intravenous antibiotics, WC wound care, AC "
                       "anticoagulation, HF heart failure). Episodes admitted before 1 July 2025 keep their program "
                       "code in the Alliance's episode records; the Alliance's stream crosswalk gives the stream each "
                       "would have under this schedule.", BODY))
    rows = [["Stream", "Name", "Principal diagnoses (ICD-10-AM categories)", "Nurse visits"]]
    for st in K.STREAMS:
        rows.append([st, K.SNAME[st], DX_TEXT[st], SCHED_TEXT.get(st, "One visit a day")])
    s.append(_table(rows, [1.4 * cm, 3.8 * cm, 6.8 * cm, 4.6 * cm]))
    s.append(Paragraph("Days at home and visits", H2))
    s.append(Paragraph("A patient's days at home run from the admission date to the separation date, both included. "
                       "The visit schedule is the minimum the program funds: the treating team may order more for a "
                       "patient, recorded as an exception in the patient's notes and not in the program's activity "
                       "returns. A visit is a nurse attending the patient's home, by a Member's nurse or by a nurse "
                       "under an Alliance agreement.", BODY))
    s.append(Paragraph("Separations", H2))
    s.append(Paragraph("A Hospital in the Home episode ends when the course is complete and the patient is discharged "
                       "(separation mode 1), when the patient returns to a ward of the Member (mode 2), is transferred "
                       "to another health service (mode 3), or dies (mode 4). A patient who returns to a ward and comes "
                       "back to the program to finish a course starts a new episode in the same stream.", BODY))
    s.append(Paragraph("Bone and joint infection (H04) is offered by Riverbend, Kingsdale and Hillcrest under the "
                       "Alliance bone and joint infection protocol. Endocarditis and other deep infection (H05) is "
                       "offered by Members with an infectious diseases service (Riverbend, Kingsdale, Hillcrest).",
                       NOTE))
    foot = _footer("%s. Hospital in the Home clinical governance." % K.ALLIANCE_SHORT, "HITH-CG-07 v2.0")
    doc = _doc(path, "Hospital in the Home program streams and visit schedule", K.PEOPLE["requester"])
    doc.build(s, onFirstPage=foot, onLaterPages=foot)


# ---------------------------------------------------------------------------------------------------- pathway

def write_pathway(path):
    import docx
    from docx.shared import Pt
    P = K.PEOPLE
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    d.add_heading("Bone and joint infection at home: Northgate Health and Westmere Health", level=1)
    d.add_paragraph("Pathway, version 1.0. Approved by the Alliance Hospital in the Home clinical governance group on "
                    "19 August 2026. Clinical lead: Dr %s, infectious diseases physician. Owners: %s (nurse unit "
                    "manager, Northgate Hospital in the Home) and %s (Hospital in the Home manager, Westmere Health)."
                    % (P["id_lead"], P["ng_num"], P["wm_mgr"]))
    d.add_heading("What changes", level=2)
    d.add_paragraph("From 1 October 2026 Northgate Health and Westmere Health admit adults with a bone and joint "
                    "infection to Hospital in the Home in stream H04, as Riverbend, Kingsdale and Hillcrest already do. "
                    "Until now these patients completed their intravenous course on a ward at Northgate and Westmere. "
                    "The pathway follows the Alliance bone and joint infection protocol (version 3, March 2025), so "
                    "assessment, antibiotic choice, monitoring and review are the same at all five Members.")
    d.add_heading("Who is eligible", level=2)
    d.add_paragraph("Adults with septic arthritis, osteomyelitis (including vertebral osteomyelitis and discitis) or a "
                    "prosthetic joint infection who have a treatment plan from the infectious diseases team, have been "
                    "clinically stable for at least 48 hours with a peripherally inserted central catheter in place, and "
                    "live within the Member's Hospital in the Home area, with a carer at home where the protocol asks for "
                    "one.")
    d.add_heading("Care at home", level=2)
    d.add_paragraph("A nurse visits every day to change the antibiotic infusion pump, check the line and take the "
                    "observations the protocol sets. Bloods are taken weekly on the visit. The infectious diseases "
                    "team reviews each patient weekly, in clinic or by telehealth. Most patients complete four to eight "
                    "weeks of intravenous therapy at home; a patient who needs a ward stay during the course comes back "
                    "to the program under a new episode to finish it.")
    d.add_heading("Admissions and activity", level=2)
    d.add_paragraph("Expected admissions from October 2026 are in each Member's 2026-27 operating plan as revised in "
                    "October 2026. Northgate and Westmere report pathway activity in the Alliance census and episode "
                    "returns in stream H04 from the first admission.")
    d.add_heading("Contacts", level=2)
    d.add_paragraph("Clinical questions: Dr %s through the infectious diseases registrar. Program questions: %s, "
                    "Alliance Hospital in the Home program manager." % (P["id_lead"], P["requester"]))
    d.core_properties.author = P["ng_num"]
    d.core_properties.title = "Bone and joint infection at home pathway"
    d.save(path)


# ---------------------------------------------------------------------------------------------------- spreadsheets

def _title_rows(ws, title, sub, note=None):
    ws.append([title])
    ws.cell(row=ws.max_row, column=1).font = Font(name="Calibri", bold=True, size=12)
    ws.append([sub])
    ws.cell(row=ws.max_row, column=1).font = Font(name="Calibri", italic=True, size=10, color="595959")
    if note:
        ws.append([note])
        ws.cell(row=ws.max_row, column=1).font = Font(name="Calibri", size=9, color="595959")
    ws.append([])


def _head(ws, cols, widths=None, fill="DDE8E2"):
    ws.append(cols)
    row = ws.max_row
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name="Calibri", bold=True)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    return row


PLAN_NOTES = {
    "NGH": "Bone and joint infection at home (stream H04) from 1 October 2026, under the Alliance pathway. Skin and "
           "soft tissue infection: emergency department direct admission extended to weekends from October 2026.",
    "RBH": "Heart failure (stream H08): admissions from the cardiology unit's new heart failure service from October "
           "2026.",
    "WMH": "Bone and joint infection at home (stream H04) from 1 October 2026, under the Alliance pathway. Urinary "
           "tract infection: emergency department direct admission from October 2026.",
    "KDH": "Respiratory infection (stream H03): the respiratory admission avoidance service moves into Hospital in "
           "the Home from October 2026.",
    "HCH": "Endocarditis and other deep infection (stream H05): admissions from the infectious diseases outreach clinic "
           "from October 2026.",
}
APPROVED = {"NGH": "23 June 2026", "RBH": "17 June 2026", "WMH": "24 June 2026", "KDH": "18 June 2026",
            "HCH": "25 June 2026"}


def write_plans(path, W):
    wb = openpyxl.Workbook()
    first = True
    for m in K.MEMBERS:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = K.SHORT[m]
        _title_rows(ws, "%s: Hospital in the Home operating plan 2026-27" % K.NAME[m],
                    "Approved by the %s executive on %s. Revision 1 of %s: activity to 30 September 2026 and changes "
                    "from October 2026." % (K.NAME[m], APPROVED[m], K.PLAN_REVISION_DATE.strftime("%-d %B %Y")),
                    PLAN_NOTES[m])
        cols = ["Stream", "Name", "Measure"] + ["%s %d%s" % (K.MON3[mo.month], mo.year % 100,
                                                            " (actual)" if mo < K.PATH else "") for mo in K.MONTHS_FY]
        cols += ["2026-27"]
        hr = _head(ws, cols, [7, 30, 13] + [10.5] * 12 + [10])
        totals = {"Admissions": [0] * 12, "Nurse visits": [0] * 12}
        for s in K.STREAMS:
            adm = [W.adm[m][s].get(mo, 0) for mo in K.MONTHS_FY]
            vis = [W.visits[m][s].get(mo, 0) for mo in K.MONTHS_FY]
            if sum(adm) == 0 and sum(vis) == 0:
                continue
            for lab, vals in (("Admissions", adm), ("Nurse visits", vis)):
                ws.append([s, K.SNAME[s], lab] + vals + [sum(vals)])
                for c in range(4, 17):
                    ws.cell(row=ws.max_row, column=c).number_format = "#,##0"
                totals[lab] = [a + b for a, b in zip(totals[lab], vals)]
        for lab in ("Admissions", "Nurse visits"):
            ws.append(["All", "All streams", lab] + totals[lab] + [sum(totals[lab])])
            for c in range(1, 17):
                ws.cell(row=ws.max_row, column=c).font = Font(name="Calibri", bold=True)
                if c >= 4:
                    ws.cell(row=ws.max_row, column=c).number_format = "#,##0"
        ws.freeze_panes = ws.cell(row=hr + 1, column=4)
        ws.append([])
        for t in ("Admissions are Hospital in the Home episodes admitted in the month; a course that resumes after a "
                  "ward stay is a new episode.",
                  "Nurse visits are the visits the program's patients at home receive on the month's days, by the "
                  "Member's nurses or under an Alliance agreement, patients admitted in earlier months included.",
                  "July to September 2026 are as recorded in the Alliance census and episode returns at 30 September "
                  "2026. October 2026 onwards are the plan.",
                  "Streams as in the Alliance program streams and visit schedule of July 2025."):
            ws.append([t])
            ws.cell(row=ws.max_row, column=1).font = Font(name="Calibri", size=9, color="595959")
    wb.save(path)


def write_capacity(path, W):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Q3 2026-27"
    _title_rows(ws, "Hospital in the Home in-house visit capacity: quarter 3 of 2026-27 (January to March 2027)",
                "Alliance Hospital in the Home program office. Version 1.1, refreshed %s from the Members' returns "
                "(version 1.0 of 9 October 2026 had Hillcrest's draft roster)." % K.CAPACITY_DATE.strftime("%-d %B %Y"),
                "In-house visits are the visits a Member's own Hospital in the Home nurses can make in the quarter as "
                "rostered. They exclude visits under the Bellbird overflow agreement and do not net any booking.")
    hr = _head(ws, ["Member", "Visiting nurse FTE", "Visiting shifts rostered", "Less leave, training and backfill",
                    "Net visiting shifts", "Visits a shift (Member standard)", "In-house visits, January to March 2027",
                    "Return received", "Returned by"], [20, 11, 12, 14, 11, 13, 15, 12, 30])
    rec = {"NGH": ("2026-10-12", "Nurse unit manager, Hospital in the Home"),
           "RBH": ("2026-10-09", "Hospital in the Home coordinator"),
           "WMH": ("2026-10-13", "Hospital in the Home manager"),
           "KDH": ("2026-10-09", "Nurse unit manager, Hospital in the Home"),
           "HCH": ("2026-10-16", "Hospital in the Home coordinator")}
    for m in K.MEMBERS:
        sh = W.shifts[m]
        fte = sh["rostered"] / 63.0 / 1.0
        ws.append([K.NAME[m], round(fte * 1.0, 1), sh["rostered"], sh["leave"], sh["net"], sh["per_shift"],
                   W.capacity[m], dt.date.fromisoformat(rec[m][0]), rec[m][1]])
        ws.cell(row=ws.max_row, column=8).number_format = "yyyy-mm-dd"
        for c in (3, 4, 5, 7):
            ws.cell(row=ws.max_row, column=c).number_format = "#,##0"
    ws.append(["All Members", None, sum(W.shifts[m]["rostered"] for m in K.MEMBERS),
               sum(W.shifts[m]["leave"] for m in K.MEMBERS), sum(W.shifts[m]["net"] for m in K.MEMBERS), None,
               sum(W.capacity[m] for m in K.MEMBERS), None, None])
    for c in range(1, 10):
        ws.cell(row=ws.max_row, column=c).font = Font(name="Calibri", bold=True)
    ws.freeze_panes = ws.cell(row=hr + 1, column=2)
    ws.append([])
    for t in ("Visiting nurse FTE: nurses rostered to home visits, full-time equivalent, averaged over the quarter's "
              "13 weeks (63 shifts a full-time nurse).",
              "Visits a shift: each Member's own standard for an eight-hour visiting shift, allowing for travel in its "
              "area.",
              "Owner: Alliance Hospital in the Home program office. Next refresh: January 2027 for April to June 2027."):
        ws.append([t])
        ws.cell(row=ws.max_row, column=1).font = Font(name="Calibri", size=9, color="595959")
    wb.save(path)


# ---------------------------------------------------------------------------------------------------- register

def write_register(path, W):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Intake log 2026"
    _title_rows(ws, "%s: intake line log, 2026" % K.PROVIDER_SHORT,
                "Calls logged by the intake desk, 2 January to 30 October 2026 (extract of 30 October 2026). Times are "
                "Melbourne time.")
    hr = _head(ws, ["Log no.", "Date", "Time", "Organisation", "Caller", "Call type", "Patient", "Notes", "Outcome",
                    "Outcome time", "Taken by"], [11, 11, 7, 30, 26, 17, 13, 50, 20, 11, 8])
    for x in W.reg:
        ws.append([x["log"], x["t"].date(), x["t"].strftime("%H:%M"), x["org"], x["caller"], x["kind"], x["patient"],
                   x["note"], x["outcome"], x["out_t"].strftime("%H:%M") if x["out_t"] is not None else None, x["by"]])
        ws.cell(row=ws.max_row, column=2).number_format = "dd/mm/yyyy"
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    ws.auto_filter.ref = "A%d:K%d" % (hr, ws.max_row)
    wb.save(path)


# ---------------------------------------------------------------------------------------------------- dictionary

DICT = [
    (K.EPISODES, "(file)", "One row per Hospital in the Home episode admitted from 1 April 2024 to 30 September 2026, "
     "from the Members' admitted episode returns to the Alliance warehouse. A patient may have several episodes."),
    (K.EPISODES, "member_code", "Alliance Member code: NGH Northgate Health, RBH Riverbend Health, WMH Westmere Health, "
     "KDH Kingsdale Health, HCH Hillcrest Health."),
    (K.EPISODES, "campus_code", "Campus of the admitting Member, as the Department's campus list gives it on the "
     "admission date (see alliance_member_campuses.csv)."),
    (K.EPISODES, "patient_no", "The Member's own patient number, in the Member's format. Numbers are unique within a "
     "Member, not across Members."),
    (K.EPISODES, "episode_no", "Episode number, unique within the extract."),
    (K.EPISODES, "admission_date", "Date the patient was admitted to Hospital in the Home (the first day at home)."),
    (K.EPISODES, "separation_date", "Date the episode ended (the last day at home). Blank for episodes still admitted "
     "at the end of 30 September 2026."),
    (K.EPISODES, "separation_mode", "1 discharged, course complete; 2 returned to a ward of the Member; 3 transferred "
     "to another health service; 4 died. Blank while the episode is open."),
    (K.EPISODES, "program_stream", "Program at admission: AB, WC, AC or HF for admissions to 30 June 2025, H01 to H08 "
     "from 1 July 2025 (program streams and visit schedule; the crosswalk maps the earlier codes)."),
    (K.EPISODES, "principal_diagnosis", "Principal diagnosis of the episode, ICD-10-AM."),
    (K.CENSUS, "(file)", "One row per patient per day at home, 1 July 2025 to 30 September 2026, from the Alliance "
     "census warehouse (loaded nightly from the Members' Hospital in the Home systems)."),
    (K.CENSUS, "program_stream", "Stream of the episode under the July 2025 streams (episodes admitted before 1 July "
     "2025 carried at their mapped stream)."),
    (K.CENSUS, "day_of_stay", "Day of the episode, 1 on the admission date."),
    (K.CENSUS, "nurse_visits", "Nurse visits the patient received that day, by the Member's nurses or under an "
     "Alliance agreement. Zero on a wound care patient's days between visits."),
    (K.CROSSWALK, "(file)", "Maps each program code used to 30 June 2025 to the July 2025 stream, by the episode's "
     "principal diagnosis category (the first three characters of the ICD-10-AM code) where the old program covers "
     "more than one stream."),
    (K.CAMPUSES, "(file)", "Campuses of the Members with Hospital in the Home episodes, with the dates each code was "
     "valid. Westmere Health's codes were reissued on 1 March 2025 when Altona Park Health merged into it."),
    (K.PLANS, "(file)", "One sheet per Member: admissions and nurse visits by stream and month, July 2026 to June "
     "2027; July to September 2026 as recorded."),
    (K.CAPACITY, "(file)", "Each Member's in-house visit capacity for January to March 2027 as rostered."),
    (K.BB_VISITS, "(file)", "Bellbird's visit records for the Alliance's patients for visits from 1 July to 30 "
     "September 2026 (Melbourne dates), as posted by its nurses' mobile app, one record a line. Extract of 5 October "
     "2026 from Bellbird's platform."),
    (K.BB_VISITS, "visit_no", "Bellbird's record number, assigned when the app posts the record."),
    (K.BB_VISITS, "client_code", "Bellbird's client account (bellbird_client_accounts.csv)."),
    (K.BB_VISITS, "patient_ref", "Bellbird's patient reference: client account, a slash, and the digits of the "
     "Member's patient number without leading zeros."),
    (K.BB_VISITS, "visit_type", "Kind of visit for the fee: IVT intravenous therapy, WND wound care, CLN clinical "
     "care."),
    (K.BB_VISITS, "start_time", "Start of the visit as the app records it, UTC."),
    (K.BB_VISITS, "end_time", "End of the visit as the app records it, UTC. Null for a cancelled visit."),
    (K.BB_VISITS, "status", "Outcome of the visit. App versions before 5.0: COMPLETE, NOTHOME (nurse attended, "
     "patient not at home or visit declined), CANCELLED. From version 5.0 (released 3 August 2026): Completed, "
     "Attempted, Cancelled."),
    (K.BB_VISITS, "recorded_at", "When the record reached Bellbird's platform, UTC. Records made without signal post "
     "when the device reconnects."),
    (K.PORTAL, "(file)", "Referrals the Members sent to Bellbird through the Alliance referral portal, sent 1 July "
     "to 30 September 2026. Extract of 2 October 2026."),
    (K.PORTAL, "sent_at", "When the Member sent the referral, Melbourne time."),
    (K.PORTAL, "requested_start", "Date from which the Member asks Bellbird to take the patient's visits."),
    (K.PORTAL, "status", "Bellbird's response: Accepted or Declined."),
    (K.PORTAL, "bellbird_response_at", "Time of Bellbird's response as returned by Bellbird's platform, UTC."),
    (K.REGISTER, "(file)", "Bellbird's intake desk log of calls to its intake line, 2026, as Bellbird sent it."),
    (K.ACCOUNTS, "(file)", "Bellbird's client accounts for the Alliance and its Members."),
    (K.CALENDAR, "(file)", "Alliance finance calendar for 2026-27: every day with its pay week (weeks run Monday to "
     "Sunday), its finance period and the Victorian public holidays."),
    (K.DECLINED, "(file)", "Referrals Bellbird declined from 1 July 2025 to 30 June 2026 with Bellbird's reason code, "
     "as Bellbird sent them with its 2025-26 service report."),
    (K.MANIFEST, "(file)", "Files in this set, with source, period and preparation."),
]


def write_dictionary(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fields"
    _title_rows(ws, "Data dictionary: Hospital in the Home overflow files", "%s, data office. 2 November 2026. "
                "Fields whose meaning the header does not give." % K.ALLIANCE_SHORT)
    _head(ws, ["File", "Field", "Definition"], [50, 22, 100])
    for f, c, d in DICT:
        ws.append([f, c, d])
        ws.cell(row=ws.max_row, column=3).alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


# ---------------------------------------------------------------------------------------------------- distractors

def write_survey(path):
    r = K.rng("survey")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "By member and quarter"
    _title_rows(ws, "Hospital in the Home patient experience survey, 2025-26",
                "Alliance quality team. Survey sent by text message to every patient discharged from Hospital in the "
                "Home, results by Member and quarter of discharge. Published 21 August 2026.")
    _head(ws, ["Member", "Quarter", "Surveys sent", "Responses", "Response rate (%)", "Care rated very good or good (%)",
               "Felt involved in decisions (%)", "Nurse came within the agreed window (%)",
               "Would recommend Hospital in the Home (%)"], [20, 14, 11, 11, 11, 13, 13, 14, 14])
    size = {"NGH": 205, "RBH": 176, "WMH": 140, "KDH": 120, "HCH": 98}
    for m in K.MEMBERS:
        for qi, q in enumerate(["Jul-Sep 2025", "Oct-Dec 2025", "Jan-Mar 2026", "Apr-Jun 2026"]):
            sent = int(size[m] * r.uniform(0.94, 1.08))
            resp = int(sent * r.uniform(0.27, 0.36))
            vals = [round(float(r.uniform(lo, hi)), 1) for lo, hi in ((92, 98), (84, 93), (81, 92), (93, 99))]
            ws.append([K.NAME[m], q, sent, resp, round(100 * resp / sent, 1)] + vals)
    wb.save(path)


def write_bb_report(path):
    r = K.rng("bb_report")
    s = [Paragraph("%s  |  Health services" % K.PROVIDER, H1S), Paragraph("Service report to Ironbark Health Alliance, "
                                                                          "2025-26", H1),
         Paragraph("Overflow nursing for the Alliance's Hospital in the Home programs, 1 July 2025 to 30 June 2026. "
                   "Prepared by %s, account manager, 28 July 2026." % K.PEOPLE["bellbird"], H1S)]
    s.append(Paragraph("Visits", H2))
    s.append(Paragraph("Bellbird's nurses made the visits below for the Alliance's patients in 2025-26, by Member and "
                       "quarter (completed visits).", BODY))
    rows = [["Member", "Jul-Sep 2025", "Oct-Dec 2025", "Jan-Mar 2026", "Apr-Jun 2026", "2025-26"]]
    base = {"NGH": 148, "RBH": 121, "WMH": 104, "KDH": 133, "HCH": 92}      # completed visits a quarter
    tot = [0] * 4
    for m in K.MEMBERS:
        v = [int(base[m] * r.uniform(0.85, 1.15)) for _ in range(4)]
        tot = [a + b for a, b in zip(tot, v)]
        rows.append([K.NAME[m]] + ["{:,}".format(x) for x in v] + ["{:,}".format(sum(v))])
    rows.append(["All Members"] + ["{:,}".format(x) for x in tot] + ["{:,}".format(sum(tot))])
    s.append(_table(rows, [4.2 * cm] + [2.5 * cm] * 5))
    s.append(Paragraph("Referral response", H2))
    s.append(Paragraph("Referrals received from the Members, the share Bellbird accepted or declined within two hours, "
                       "and the share of accepted referrals Bellbird first visited within 24 hours, as Bellbird's "
                       "platform reports them.", BODY))
    rows = [["Member", "Referrals", "Responded within 2 hours (%)", "First visit within 24 hours (%)"]]
    for m in K.MEMBERS:
        rows.append([K.NAME[m], "{:,}".format(int(base[m] * 1.28 * r.uniform(0.9, 1.1))),
                     "%.1f" % float(r.uniform(84, 93)), "%.1f" % float(r.uniform(86, 95))])
    s.append(_table(rows, [4.2 * cm, 2.6 * cm, 4.6 * cm, 4.6 * cm]))
    s.append(Paragraph("Workforce and quality", H2))
    s.append(Paragraph("Bellbird ran 29 registered nurses on the Alliance's work through the year, from its Coburg and "
                       "Sunshine bases. Eleven incidents were reported to the Members (none above minor harm). Seven "
                       "complaints were received, all resolved within 30 days. The mobile app moved to version 4.6 in "
                       "February 2026; version 5.0 is planned for August 2026.", BODY))
    s.append(Paragraph("Figures are from Bellbird's platform and may differ from the Members' own records.", NOTE))
    foot = _footer(K.PROVIDER, "Service report 2025-26")
    doc = _doc(path, "Service report to Ironbark Health Alliance 2025-26", K.PEOPLE["bellbird"])
    doc.build(s, onFirstPage=foot, onLaterPages=foot)
