"""Assertions on the built pack and the goldens: the call and every rung of the ladder reproduce from the shipped
files, every semantic-fork row (SF1 to SF14) is closed as the design note says, the client's three rules hold on the
clean truth (R1, R2, R3), the asks' devices and stops land where the ask sheet says, the goldens carry the computed
figures, and the input gates and the realism sweep hold. Run by build.py."""
from __future__ import annotations

import collections
import datetime as dt
import io
import itertools
import json
import math
import os
import re
import shutil
import tempfile
import zipfile

import numpy as np
import openpyxl

import calc as C
import common as K
import model as MD
import world as WD

N = [0]
GOLD = {"NGH": 300, "RBH": 250, "WMH": 225, "KDH": 75, "HCH": 0}
REQ50 = {"NGH": 2400, "RBH": 2150, "WMH": 1800, "KDH": 1400, "HCH": 1250}
DEFAULT = {"NGH": 100, "RBH": 50, "WMH": 50, "KDH": 25, "HCH": 25}


def ok(cond, msg):
    if not cond:
        raise AssertionError("CHECK FAILED: " + msg)
    N[0] += 1


def edge_near(x, step):
    """Distance from x to the nearer edge of the bin it rounds into (nearest-step rounding)."""
    q = x / step
    return (0.5 - abs(q - math.floor(q + 0.5))) * step


def residue(excess):
    """Visits of the excess inside its last lot of 25 (0 to 25)."""
    return excess - (K.lots_for(excess) - 1) * K.LOT


def fmt(bk):
    return "/".join(str(bk[m]) for m in K.MEMBERS)


def pdf_text(path):
    from pypdf import PdfReader
    return re.sub(r"\s+", " ", " ".join(p.extract_text() or "" for p in PdfReader(path).pages))


def doc_text(path):
    if path.endswith(".pdf"):
        return pdf_text(path)
    if path.endswith(".docx"):
        with zipfile.ZipFile(path) as z:
            x = z.read("word/document.xml").decode("utf-8")
        return re.sub(r"\s+", " ", " ".join(re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", x)))
    if path.endswith(".xlsx"):
        wb = openpyxl.load_workbook(path, data_only=True)
        return " ".join(str(v) for ws in wb.worksheets for r in ws.iter_rows(values_only=True) for v in r
                        if v is not None)
    with open(path, encoding="utf-8") as fh:
        return fh.read()


def run(target, golden, W=None, meta_path=None):
    B = C.load(target)
    G = C.golden(B)
    D = C.default_path(B)
    Rv = C.rivals(B, G)
    F = C.family(B, G)
    A = C.load_asks(target)
    texts = {f: doc_text(os.path.join(target, f)) for f in os.listdir(target)
             if f.endswith((".pdf", ".docx", ".xlsx"))}

    # ------------------------------------------------------------------ the call
    ok(G["booking"] == GOLD, "the golden booking %s" % fmt(G["booking"]))
    ok(G["left"] == 0 and sum(G["booking"].values()) == K.BLOCK, "the 850 visits are used")
    ok({m: K.round_to(G["reqs"][m], 50) for m in K.MEMBERS} == REQ50, "requirements to the nearest 50")
    for m in K.MEMBERS:
        ok(edge_near(G["reqs"][m], 50) >= 15, "%s requirement %.1f at least 15 visits from a 50-visit edge" %
           (m, G["reqs"][m]))
        ok(G["reqs"][m] > B.capacity[m], "SF7 %s percentile above its in-house capacity" % m)
        ok(7 <= residue(G["reqs"][m] - B.capacity[m]) <= 18, "%s excess %.1f at least 7 visits inside its lot" %
           (m, G["reqs"][m] - B.capacity[m]))
    ok(G["need"]["NGH"] == 12 and G["need"]["RBH"] == 10 and G["need"]["WMH"] == 9, "lots needed by the first three")
    ok(G["need"]["KDH"] >= 5, "Kingsdale needs %d lots, more than the 3 left" % G["need"]["KDH"])
    for m in K.MEMBERS:
        mix, mean = MD.quarter_mix(B.adm[m], G["pool"])
        pmf = MD.compound(mix)
        x = np.arange(len(pmf))
        ok(abs(pmf.sum() - 1) < 1e-9 and abs((x * pmf).sum() - mean) < 0.05, "%s compound pmf exact" % m)
        ok(pmf[-200:].sum() < 1e-12, "%s pmf grid wide enough" % m)
        ok(abs(mean - B.plan_q[m]) < 3.0, "%s the plan's visits are the model's expectation (%.1f against %d)" %
           (m, mean, B.plan_q[m]))
    ok(set(B.plan_q.values()) == {2143, 1905, 1570, 1209, 1064}, "plans' January to March 2027 visits")
    ok(B.capacity == {"NGH": 2112, "RBH": 1911, "WMH": 1589, "KDH": 1246, "HCH": 1086}, "capacity return")

    # ------------------------------------------------------------------ the ladder
    ok(D["booking"] == DEFAULT and D["left"] == 600, "default path (plan plus Poisson) books %s, %d left" %
       (fmt(D["booking"]), D["left"]))
    expect = {"R0_poisson": "100/50/50/25/25", "R0_poisson_normal": "100/50/50/25/25",
              "R0_daily_independent": "125/50/50/25/50", "R1_monthly_ep": "250/225/150/125/100",
              "R1_weekly_ep": "200/125/125/75/125", "R2_quarterly_ep": "225/325/175/125/0",
              "R1_nb_monthly_ep": "250/225/150/125/100", "R1_monthly_cen": "225/175/150/125/175",
              "R1_weekly_cen": "200/100/125/75/125", "R2_quarterly_cen": "150/125/175/200/200",
              "R1_nb_monthly_cen": "225/175/150/125/175", "R1_monthly_deseasonalised_ep": "200/175/125/75/150",
              "R2b_own_history": "175/250/100/175/150", "R3x_fixed_admissions": "150/100/75/50/50",
              "R3x_cohort": "325/275/250/0/0", "R3x_no_carry_in": "275/225/200/150/0",
              "R3x_fixed_stays": "275/225/200/150/0", "R3x_whole_stay_normal": "325/275/250/0/0",
              "R1_ar1_daily": "225/150/150/75/125"}
    ok(set(expect) == set(Rv), "every rival on the ladder: %s" % sorted(set(Rv) ^ set(expect)))
    for k, v in expect.items():
        ok(fmt(Rv[k]["booking"]) == v, "rival %s books %s (got %s)" % (k, v, fmt(Rv[k]["booking"])))
        ok(sum(Rv[k]["booking"][m] != GOLD[m] for m in K.MEMBERS) >= 4, "rival %s differs at four members or more" % k)
        ok(any(Rv[k]["booking"][m] != GOLD[m] for m in K.MEMBERS[:3]), "rival %s differs at the first three" % k)
    mg = {k: {m: Rv[k]["reqs"][m] - B.plan_q[m] for m in K.MEMBERS} for k in Rv}
    gm = {m: G["reqs"][m] - B.plan_q[m] for m in K.MEMBERS}
    ok(gm["NGH"] - mg["R0_poisson"]["NGH"] >= 150, "dominance: Northgate's margin moves %.0f from the Poisson margin" %
       (gm["NGH"] - mg["R0_poisson"]["NGH"]))
    for k in ("R1_monthly_ep", "R1_monthly_deseasonalised_ep", "R1_ar1_daily", "R1_weekly_ep"):
        ok(gm["NGH"] - mg[k]["NGH"] >= 40 and gm["WMH"] - mg[k]["WMH"] >= 40,
           "dominance: the history's spread (%s) understates Northgate and Westmere by 40 visits or more" % k)
    for m in K.MEMBERS:
        ok(gm[m] >= 3 * mg["R0_poisson"][m], "%s margin %.0f at least three Poisson margins" % (m, gm[m]))
    ok(all(K.round_to(D["reqs"][m], 50) != REQ50[m] for m in K.MEMBERS), "default requirements miss every 50-bin")
    # the near rivals flank the golden: one lot below (stays not varied, no carry-in) and above (whole stays)
    for m in K.MEMBERS[:3]:
        ok(K.lots_for(Rv["R3x_fixed_stays"]["reqs"][m] - B.capacity[m]) == G["need"][m] - 1,
           "%s: stays given their mean books one lot fewer" % m)
        ok(K.lots_for(Rv["R3x_whole_stay_normal"]["reqs"][m] - B.capacity[m]) == G["need"][m] + 1,
           "%s: the quarter's admissions with whole stays books one lot more" % m)

    # ------------------------------------------------------------------ the golden family
    ok(len(F) == 13, "thirteen admissible constructions")
    res_min = 25.0
    for k, v in F.items():
        ok(v["booking"] == GOLD, "family %s books the golden (%s)" % (k, fmt(v["booking"])))
        ok({m: K.round_to(v["reqs"][m], 50) for m in K.MEMBERS} == REQ50, "family %s requirements" % k)
        for m in K.MEMBERS[:3]:
            r = residue(v["reqs"][m] - B.capacity[m])
            res_min = min(res_min, r, 25 - r)
        spread = max(abs(v["reqs"][m] - G["reqs"][m]) for m in K.MEMBERS[:3])
        ok(spread <= 8.0, "family %s within 8 visits of the golden at the first three (%.1f)" % (k, spread))
    ok(res_min >= 5.5, "every admissible construction at least 5.5 visits inside the golden lot (%.1f)" % res_min)
    for seed in (3, 5, 8, 13, 21):
        sim = {m: C.simulate_p90(B, m, G["pool"], 20000, seed=seed)[0] for m in K.MEMBERS}
        ok(MD.book(sim, B.capacity)[0] == GOLD, "a 20,000-quarter simulation (seed %d) books the golden" % seed)
    disp = C.dispersion(B)
    ok(all(0.9 <= v <= 1.05 for v in disp.values()), "daily admissions Poisson around each month's mean %s" %
       {m: round(v, 3) for m, v in disp.items()})

    # ------------------------------------------------------------------ the semantic fork table
    _forks(target, B, G, Rv, texts)
    # ------------------------------------------------------------------ the three rules (R1, R2, R3)
    _three_rules(target, B, G, Rv, W, texts)
    # ------------------------------------------------------------------ the asks
    asks = _asks(target, A, W)
    # ------------------------------------------------------------------ the goldens
    _goldens(golden, B, G, asks)
    # ------------------------------------------------------------------ input gates, distractors, realism
    _gates(target, golden, B, G, A, asks, W, texts, meta_path)
    print("checks: %d assertions passed" % N[0])
    return N[0]


# ---------------------------------------------------------------------------------------------------- SF rows

def _forks(target, B, G, Rv, texts):
    s2 = texts[K.SCHED2]
    # SF1: the percentile of the quarter taken together, not daily or weekly percentiles scaled up
    ok("the visits of the Quarter taken together" in s2, "SF1 clause 2.4 files the quarter taken together")
    ep, cen = C.history_daily(B)
    p_day = {m: float(np.percentile(cen[m], 90)) * 90 / float(cen[m].mean()) * B.plan_q[m] / 90 for m in K.MEMBERS}
    bk = MD.book(p_day, B.capacity)[0]
    ok(bk != GOLD and bk["HCH"] == 0 and sum(bk.values()) == K.BLOCK and bk["KDH"] == 0,
       "SF1 the daily percentile times the days uses the block by Westmere (%s)" % fmt(bk))
    # SF2: every visit, whoever makes them
    ok("whoever makes them" in s2, "SF2 clause 2.3 counts every visit, whoever makes them")
    ok(any("under an Alliance agreement" in t for t in [texts[K.PLANS]]), "SF2 the plan's visits include agreements")
    # SF2 (continued): visits the treating team orders beyond the schedule are outside the activity returns; at three
    # times the rate Bellbird's own records show for intravenous patients they leave the booking where it is
    iv_days = collections.Counter()
    extra = 0
    vs = C.dedupe(C.load_asks(target).visits)
    by_day = collections.defaultdict(list)
    for v in vs:
        if v["status"] != "cancel" and v["type"] == "IVT":
            by_day[(v["pref"], v["start"].date())].append(v)
    for k, lst in by_day.items():
        iv_days[k] += 1
        done = [v for v in lst if v["status"] == "done"]
        extra += max(0, len(done) - 1)
    rate = extra / max(1, len(iv_days))
    ok(0.01 < rate < 0.08, "SF2 unscheduled same-day returns on %.3f of Bellbird's intravenous patient-days" % rate)
    infl = {}
    for m in K.MEMBERS:
        iv = sum(B.vis[m][s].get(dt.date(2027, mo, 1), 0) for s in K.INFECTION for mo in (1, 2, 3))
        sd = G["det"][m]["sd"]
        infl[m] = G["reqs"][m] + 1.2815516 * (math.sqrt(sd * sd + 3 * rate * iv) - sd)
    ok(MD.book(infl, B.capacity)[0] == GOLD, "SF2 exception visits at three times that rate: the same booking")
    # SF3: the quarter's days, patients admitted before it included
    ok("with patients admitted before the Quarter included for their days inside it" in s2, "SF3 clause 2.3 window")
    ok(Rv["R3x_cohort"]["booking"] != GOLD and Rv["R3x_no_carry_in"]["booking"] != GOLD,
       "SF3 the cohort and no-carry-in readings book otherwise")
    # SF4: the plan is the expectation
    ok("operating plan as revised in October 2026" in s2, "SF4 clause 2.5 files the plan as the expectation")
    nr = {m: C.requirement(B, m, recentre=False)[0] for m in K.MEMBERS}
    ok(MD.book(nr, B.capacity)[0] == GOLD, "SF4 recentring on the plan or not, the same booking")
    # SF5: in-house capacity is the return's figure, net shifts times visits a shift
    wb = openpyxl.load_workbook(os.path.join(target, K.CAPACITY), data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r[0] == "Member")
    for r in rows[hi + 1:hi + 6]:
        ok(r[4] * r[5] == r[6] and r[2] - r[3] == r[4], "SF5 %s in-house visits are net shifts times the standard" % r[0])
    # SF6: lots that cover the excess (round up); nearest-lot rounding books otherwise
    ok("the lots of 25 visits that cover the amount" in s2, "SF6 clause 2.6 rounds up")
    near = {}
    left = K.BLOCK // K.LOT
    for m in K.MEMBERS:
        n = max(0, int(math.floor((G["reqs"][m] - B.capacity[m]) / K.LOT + 0.5)))
        near[m] = min(n, left) * K.LOT
        left -= min(n, left)
    ok(near != GOLD, "SF6 nearest-lot rounding books %s" % fmt(near))
    # SF7 asserted with the call (no member under its capacity); the clause exists
    ok("does not exceed its in-house capacity books no lots" in s2, "SF7 clause 2.6")
    # SF8: in order until the block is used; pro rata or skipping books otherwise
    ok("takes the lots that remain, and the Members after it book none" in s2, "SF8 clause 2.7")
    need = G["need"]
    tot = sum(need.values())
    pro = {m: int(math.floor(need[m] * (K.BLOCK // K.LOT) / tot)) * K.LOT for m in K.MEMBERS}
    ok(pro != GOLD, "SF8 pro rata books %s" % fmt(pro))
    # SF9: the listed order
    ok("Northgate Health, Riverbend Health, Westmere Health, Kingsdale Health, Hillcrest Health" in s2, "SF9 order")
    by_size = sorted(K.MEMBERS, key=lambda m: -B.plan_q[m])
    ok(by_size == K.MEMBERS, "SF9 the order by size is the listed order (same call)")
    by_excess = sorted(K.MEMBERS, key=lambda m: -(G["reqs"][m] - B.capacity[m]))
    ok(MD.book(G["reqs"], B.capacity, order=by_excess)[0] != GOLD, "SF9 an order by excess books otherwise")
    # SF10: the pathway members take the Alliance's bone and joint stays; per-member stays elsewhere converge
    pw = texts[K.PATHWAY]
    ok("follows the Alliance bone and joint infection protocol" in pw, "SF10 the pathway follows the Alliance protocol")
    for m in K.PATHWAY_MEMBERS:
        ok(len(C.stays(B, member=m)["H04"]) == 0, "SF10 %s has no bone and joint stays of its own" % m)
    # SF11, SF12, SF13 converge (the family carries them); SF14 the requirement's bin (asserted with the call)
    ok(sum(1 for e in B.eps if e["sep"] is None) >= 30, "SF11 open episodes on 30 September 2026 exist")
    ok(sum(1 for e in B.eps if e["mode"] == "2") > 100, "SF12 ward returns exist")
    ok(sum(1 for e in B.eps if e["code"] in K.OLD_PROGRAMS) > 2500, "SF13 old program codes exist")
    ok("nearest 50 visits" in open(os.path.join(os.path.dirname(target), "prompt.txt")).read(),
       "SF14 the prompt sets the requirement's rounding")


# ---------------------------------------------------------------------------------------------------- R1 to R3

def _three_rules(target, B, G, Rv, W, texts):
    # R1: the raw material ships (episodes with admission and separation dates; the plan's admissions by stream and
    # month); no shipped column or document holds a percentile, a margin, a requirement or a booking
    ok(len(B.eps) >= 7000 and sum(1 for e in B.eps if e["sep"]) >= 6500, "R1 %d episodes with their stays" % len(B.eps))
    ok(all(len(G["pool"][s]) >= 60 for s in ("H04", "H05")), "R1 the long-stay streams have %s completed stays" %
       {s: len(G["pool"][s]) for s in ("H04", "H05")})
    ok(sum(len(v) for v in G["pool"].values()) >= 5000, "R1 the stays pool reads thousands of completed episodes")
    ok(all(sum(B.adm[m]["H04"].get(dt.date(2026, mo, 1), 0) for mo in (10, 11, 12)) > 0 for m in K.PATHWAY_MEMBERS),
       "R1 the plan carries the pathway members' bone and joint admissions")
    for f in os.listdir(target):
        if f.endswith(".csv"):
            with open(os.path.join(target, f), encoding="utf-8") as fh:
                head = fh.readline().lower()
            ok(not re.search(r"percentile|p90|margin|requirement|lots?\b|booking", head), "R1 %s has no answer column" % f)
    for f, t in texts.items():
        for v in (G["booking"]["NGH"], G["booking"]["RBH"], G["booking"]["WMH"]):
            pass
        ok(not re.search(r"\b(2,?397|2,?148|1,?798|1,?403|1,?247)\b", t), "R1 %s states no requirement" % f)
        if f != K.SCHED2:
            ok("90th percentile" not in t and "percentile" not in t.lower(), "R1 %s names no percentile" % f)
        ok(not re.search(r"poisson|overdispers|negative binomial|simulat|monte carlo", t, re.I),
           "R1 %s names no estimator" % f)
    # R2: every Layer 1 item mishandled the natural way leaves the call where it was
    pool_clean = WD.pools(W.eps)
    ok(all(sorted(pool_clean[s].tolist()) == sorted(G["pool"][s].tolist()) for s in K.STREAMS),
       "R2 the stays read through the crosswalk equal the clean truth's")
    post = C.stays(B, first=K.REVISION)
    ok(MD.book({m: C.requirement(B, m, post)[0] for m in K.MEMBERS}, B.capacity)[0] == GOLD,
       "R2 old program codes dropped instead of mapped: the same booking")
    valid_now = {"4103", "4211", "4218", "5140", "5141", "5212", "5307"}
    eps_save = B.eps
    B.eps = [e for e in eps_save if e["campus"] in valid_now]
    ok(len(B.eps) < len(eps_save), "R2 Westmere's reissued campus codes are present")
    cur = C.stays(B)
    B.eps = eps_save
    ok(MD.book({m: C.requirement(B, m, cur)[0] for m in K.MEMBERS}, B.capacity)[0] == GOLD,
       "R2 episodes under Westmere's retired campus codes dropped: the same booking")
    hf_as_code = {k: (v if k != "H08" else np.array([(e["sep"] - e["adm"]).days + 1 for e in B.eps
                                                     if e["code"] == "H08" and e["sep"] and e["adm"] <= K.POOL_LAST_ADMISSION]))
                  for k, v in G["pool"].items()}
    ok(len(hf_as_code["H08"]) < len(G["pool"]["H08"]) and
       MD.book({m: C.requirement(B, m, hf_as_code)[0] for m in K.MEMBERS}, B.capacity)[0] == GOLD,
       "R2 the old heart failure code HF read as a stream of its own (its stays lost): the same booking")
    merged = C.stays(B, merge=True)
    ok(MD.book({m: C.requirement(B, m, merged)[0] for m in K.MEMBERS}, B.capacity)[0] == GOLD,
       "R2 a course split by a ward return joined into one stay: the same booking")
    ok(MD.book({m: C.requirement(B, m, C.stays(B, include_open=True, include_late=True))[0] for m in K.MEMBERS},
               B.capacity)[0] == GOLD, "R2 open episodes taken at their stay to date: the same booking")
    wb = openpyxl.load_workbook(os.path.join(target, K.PLANS), data_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        hi = next(i for i, r in enumerate(rows) if r[0] == "Stream")
        cols = [j for j, h in enumerate(rows[hi]) if h in ("Jan 27", "Feb 27", "Mar 27")]
        allrow = next(r for r in rows if r[0] == "All" and r[2] == "Nurse visits")
        m = {v: k for k, v in K.SHORT.items()}[ws.title]
        ok(sum(allrow[j] for j in cols) == B.plan_q[m], "R2 %s plan's all-streams row gives the same expectation" % m)
    # R3a: the default path on the clean truth files the named wrong booking
    from scipy import stats
    r0 = {m: float(stats.poisson.ppf(0.9, W.plan_q[m])) for m in K.MEMBERS}
    bk0, _, left0 = MD.book(r0, W.capacity)
    ok(bk0 == DEFAULT and left0 == 600, "R3a the default path on the clean truth books %s, %d left" % (fmt(bk0), left0))
    # R3b: only the decisive step changed (the quarter built from admissions and stays) files the golden; each
    # lower rung's single correction stays wrong
    r3 = {m: MD.requirement(W.adm[m], pool_clean, W.plan_q[m])[0] for m in K.MEMBERS}
    ok(MD.book(r3, W.capacity)[0] == GOLD, "R3b the decisive step alone on the clean truth books the golden")
    days = K.days(dt.date(2024, 7, 1), K.D1)
    di = {d: i for i, d in enumerate(days)}
    Hc = {m: np.zeros(len(days)) for m in K.MEMBERS}
    for e in W.eps:
        for k in range(e["stay"]):
            d = e["adm"] + dt.timedelta(days=k)
            if d in di:
                Hc[e["member"]][di[d]] += K.visits_on_day(e["stream"], k)
    z = 1.2815516
    mon = {}
    for m in K.MEMBERS:
        mo = collections.OrderedDict()
        for d, v in zip(days, Hc[m]):
            mo[(d.year, d.month)] = mo.get((d.year, d.month), 0) + v
        vals = np.array(list(mo.values()))
        mon[m] = W.plan_q[m] + z * math.sqrt(3) * vals.std(ddof=1) / vals.mean() * W.plan_q[m] / 3
    ok(MD.book(mon, W.capacity)[0] != GOLD, "R3b the history's monthly spread alone stays wrong")
    fixed = Rv["R3x_fixed_stays"]["booking"]
    ok(fixed != GOLD, "R3b admissions without varied stays stay wrong")
    # closure axes 14 and 15: row order does not matter; one row per episode
    eps_save = B.eps
    B.eps = list(reversed(eps_save))
    rev = C.stays(B)
    B.eps = eps_save
    ok(all(sorted(rev[s_].tolist()) == sorted(G["pool"][s_].tolist()) for s_ in K.STREAMS), "axis 14 row order inert")
    ok(len({e["episode"] for e in B.eps}) == len(B.eps), "axis 15 one row per episode")
    # R3c: no plant: no copied block in any Member's daily visits, no duplicated census day, no document with a booking
    for m in K.MEMBERS:
        s = Hc[m]
        blocks = collections.Counter(tuple(s[i:i + 14]) for i in range(0, len(s) - 14))
        ok(max(blocks.values()) == 1, "R3c no repeated 14-day block in %s's daily visits" % m)
    seen = set()
    dup = 0
    with open(os.path.join(target, K.CENSUS), encoding="utf-8") as fh:
        next(fh)
        for line in fh:
            p = line.split(",")
            key = (p[0], p[4])
            dup += key in seen
            seen.add(key)
    ok(dup == 0, "R3c one census row per episode and day")
    for f, t in texts.items():
        ok(not re.search(r"\b(300|250|225)\s+visits\b.*\b(book|lot)", t, re.I) or f == K.SCHED2,
           "R3c %s carries no booking" % f)


# ---------------------------------------------------------------------------------------------------- the asks

def _count(k, word):
    return "%d %s%s" % (k, word, "" if k == 1 else "s")


def _xlsx_size(path):
    """A workbook's size as a manifest gives it: its sheets when it has several, else the rows of its table (below
    the header, down to the first totals or note row)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    if len(wb.worksheets) > 1:
        return _count(len(wb.worksheets), "sheet")
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if sum(v is not None for v in r) >= 3)
    n = 0
    for r in rows[hi + 1:]:
        if sum(v is not None for v in r) < 2 or str(r[0]).startswith("All "):
            break
        n += 1
    return _count(n, "row")


def _round(x):
    return int(math.floor(x + 0.5))


def _r1(x):
    return math.floor(x * 10 + 0.5) / 10


def _asks(target, A, W):
    a = C.ask_a(A)
    vs = C.dedupe(A.visits)
    n_rep = sum(1 for x in W.bb if x["replay"])
    ok(len(A.visits) - len(vs) == n_rep, "A the overlap test removes exactly the %d re-posted records" % n_rep)
    for m in K.MEMBERS:
        for mo in (7, 8, 9):
            ok(any(x["replay"] and x["member"] == m and x["start"].month == mo for x in W.bb),
               "A a re-posted record in %s %d" % (m, mo))
    # a re-posted record and the record it repeats: one Melbourne day, one side of 7.00 am and 6.00 pm, of the
    # referral's acceptance and of 24 hours after it, starts 0 to 12 minutes apart, the same status and kind of visit
    pairs = [(x, x["orig"]) for x in W.bb if x["replay"]]
    ok(len(pairs) == n_rep and all(WD.lines(x["start"], o) == WD.lines(o["start"], o) for x, o in pairs),
       "A every re-post pair on one side of every line its fee, day and clocks turn on (%d pairs)" % len(pairs))
    ok(all(abs((x["start"] - o["start"]).total_seconds()) <= 12 * 60 and x["status"] == o["status"] and
           x["type"] == o["type"] and x["client"] == o["client"] and x["pref"] == o["pref"] for x, o in pairs),
       "A re-post pairs start 0 to 12 minutes apart, one status, kind, account and reference")
    by_pair = collections.defaultdict(list)
    for v in A.visits:
        if v["status"] != "cancel":
            by_pair[(v["nurse"], v["pref"])].append(v)
    olap = [(v, w) for lst in by_pair.values() for v, w in itertools.combinations(lst, 2)
            if v["start"] < w["end"] and w["start"] < v["end"]]
    ok(len(olap) == n_rep and all(v["start"].date() == w["start"].date() and
                                  C.fee(v, A.holidays) == C.fee(w, A.holidays) for v, w in olap),
       "A on the export: %d overlapping pairs, each on one Melbourne day at one fee" % len(olap))
    # amounts: exact cents, never on a half-dollar tie
    for k, v in a.items():
        ok(abs(v * 100 - round(v * 100)) < 1e-6 and abs((v - math.floor(v)) - 0.5) > 0.004,
           "A %s %d amount %.2f in whole cents, not on a tie" % (k[0], k[1], v))
    # the stops: every one lands off every cell
    raw = C.ask_a(A, visits=A.visits)
    ok(all(_round(raw[k]) != _round(a[k]) for k in a), "A natural read (re-posts kept) misses every cell")
    utc = []
    for v in vs:
        w = dict(v)
        w["start"] = dt.datetime.fromisoformat(v["raw"]["start_time"])
        utc.append(w)
    ok(all(_round(C.ask_a(A, visits=utc)[k]) != _round(a[k]) for k in a), "A UTC read as Melbourne misses every cell")
    utc_raw = []
    for v in A.visits:
        w = dict(v)
        w["start"] = dt.datetime.fromisoformat(v["raw"]["start_time"])
        utc_raw.append(w)
    ok(all(_round(C.ask_a(A, visits=utc_raw)[k]) != _round(a[k]) for k in a), "A exact-duplicate sweep, UTC kept")
    seen, spd = set(), []
    for v in sorted(A.visits, key=lambda v: (v["rec"], v["no"])):
        key = (v["pref"], v["start"].date())
        if v["status"] != "cancel" and key in seen:
            continue
        seen.add(key)
        spd.append(v)
    sp = C.ask_a(A, visits=spd)
    ok(sum(_round(sp[k]) != _round(a[k]) for k in a) >= 13, "A the same-patient-same-day sweep misses %d cells" %
       sum(_round(sp[k]) != _round(a[k]) for k in a))
    nov = [v for v in vs if v["raw"]["status"] in ("Completed", "Attempted", "Cancelled")]
    nv = C.ask_a(A, visits=nov)
    ok(all(_round(nv[(m, 7)]) != _round(a[(m, 7)]) for m in K.MEMBERS), "A the old status codes dropped: July misses")
    noalt = [v for v in vs if v["client"] != WD.CLIENT_ALTONA]
    na = C.ask_a(A, visits=noalt)
    ok(all(_round(na[("WMH", mo)]) != _round(a[("WMH", mo)]) for mo in (7, 8, 9)),
       "A Altona Park's account dropped: Westmere misses every month")
    off = [v for v in vs if v["client"] != WD.CLIENT_OFFICE]
    no = C.ask_a(A, visits=off)
    ok(all(_round(no[(m, 8)]) != _round(a[(m, 8)]) for m in K.MEMBERS),
       "A the billing account's intake-file records dropped: every Member's August misses")
    ok(all(v["client"] == WD.CLIENT_OFFICE for v in A.visits if "/IL26-" in v["pref"]) and
       all("/IL26-" in v["pref"] for v in A.visits if v["client"] == WD.CLIENT_OFFICE), "A intake files, and only "
                                                                                     "they, on the billing account")
    nh = C.ask_a(A, holidays=set())
    ok(sum(_round(nh[(m, 9)]) != _round(a[(m, 9)]) for m in K.MEMBERS) >= 3, "A the holiday unloaded moves September")
    # the chart: payable visits by week; re-posts kept moves every Member's total
    wk = C.weekly(A)
    wr = C.weekly(A, visits=A.visits)
    tot = {m: sum(wk[(m, w)] for w in C.WEEKS) for m in K.MEMBERS}
    ok(all(sum(wr[(m, w)] for w in C.WEEKS) != tot[m] for m in K.MEMBERS), "chart re-posts kept move every total")
    wo = C.weekly(A, visits=off)
    ok(all(sum(wo[(m, w)] for w in C.WEEKS) != tot[m] for m in K.MEMBERS), "chart intake files dropped move every "
                                                                            "total")
    ok(len(set(tot.values())) == 5, "chart totals distinct")
    cal = C._csv(os.path.join(target, K.CALENDAR))
    ok({r["week_commencing"] for r in cal if K.Q3S.isoformat() <= r["date"] <= K.Q3E.isoformat()} ==
       {w.isoformat() for w in C.WEEKS}, "chart the calendar's pay weeks are the chart's weeks")
    # Ask B
    b = C.ask_b(A)
    po = C.ask_b(A, refs=C.referrals(A, register=False))
    nl = C.ask_b(A, link=False)
    for m in K.MEMBERS:
        x = b[m]
        ok(po[m]["n"] != x["n"] and _r1(po[m]["pct_acc2"]) != _r1(x["pct_acc2"]) and
           _r1(po[m]["pct_v24"]) != _r1(x["pct_v24"]), "B %s portal only misses all three figures" % m)
        ok(_r1(nl[m]["pct_v24"]) != _r1(x["pct_v24"]), "B %s telephoned referrals found, intake files unlinked: v24 "
                                                       "misses" % m)
        for k, n in ((x["acc2"], x["n"]), (x["v24"], x["acc"])):
            v = 1000.0 * k / n
            ok(abs(v - math.floor(v) - 0.5) > 1e-6, "B %s share %d/%d not on a tie" % (m, k, n))
        ok(min(abs(l - 24) for l in x["lags"]) >= 4 / 60 and min(abs(l - 120) for l in x["resp_lags"]) >= 2,
           "B %s no first-visit lag within 4 minutes of 24 hours, no response within 2 minutes of 2 hours" % m)
        ok(len(x["lags"]) == x["acc"], "B %s every accepted referral has a first visit on file" % m)
    allrows = []
    for r in A.register:
        if r["Organisation"] in C.MEMBER_BY_ORG and r["Call type"] in ("Referral", "Referral follow-up"):
            d = r["Date"].date() if isinstance(r["Date"], dt.datetime) else r["Date"]
            if K.Q3S <= d <= K.Q3E:
                allrows.append(C.MEMBER_BY_ORG[r["Organisation"]])
    cnt = collections.Counter(allrows)
    ok(sum(po[m]["n"] + cnt[m] != b[m]["n"] for m in K.MEMBERS) >= 4, "B follow-up calls counted as referrals miss")
    utc_b = C.referrals(A)
    for x in utc_b:
        if x["source"] == "portal":
            x["resp"] = x["resp"] - dt.timedelta(hours=10)
    ub = C.ask_b(A, refs=utc_b)
    ok(sum(_r1(ub[m]["pct_acc2"]) != _r1(b[m]["pct_acc2"]) for m in K.MEMBERS) == 5, "B UTC response kept misses acc2")
    reps = C.ask_b(A, visits=A.visits)
    ok(all(_r1(reps[m]["pct_v24"]) == _r1(b[m]["pct_v24"]) for m in K.MEMBERS), "B re-posts do not move v24 (same)")
    nph = {m: sum(1 for x in W.refs if x["phone"] and x["member"] == m and not x["declined"]) for m in K.MEMBERS}
    ok(min(nph.values()) >= 1, "B every Member telephoned an accepted referral in the outage %s" % nph)
    ok(all(not (K.OUTAGE[0] <= dt.datetime.strptime(r["sent_at"], "%Y-%m-%d %H:%M") < K.OUTAGE[1]) for r in A.portal),
       "B the portal holds nothing sent during the outage")
    ok(sum(1 for v in A.visits if "/IL26-" in v["pref"]) >= 10, "B intake-file references on the visit export")
    ok(all(v["pref"].split("/", 1)[1] in A.intake for v in A.visits if "/IL26-" in v["pref"]),
       "B every intake-file reference is a telephoned referral's log number")
    # whichever record of a re-post pair is kept (first or last posted, lowest or highest visit number, earliest or
    # latest start), every amount, every pay week's count and every referral figure is the same
    sig_b = lambda bb: {m: (bb[m]["n"], bb[m]["acc2"], bb[m]["acc"], bb[m]["v24"]) for m in K.MEMBERS}
    for rule in C.KEEP:
        vk = C.dedupe(A.visits, keep=rule)
        ok(len(A.visits) - len(vk) == n_rep, "survivor %s: the overlap test removes the %d re-posts" % (rule, n_rep))
        ak = C.ask_a(A, visits=vk)
        ok(all(round(ak[k], 2) == round(a[k], 2) for k in a), "survivor %s: all 15 amounts to the cent" % rule)
        ok(C.weekly(A, visits=vk) == wk, "survivor %s: every Member's count in every pay week" % rule)
        ok(sig_b(C.ask_b(A, visits=vk)) == sig_b(b), "survivor %s: every referral count and share" % rule)
    # the referral response's bases, as Schedule 4 clause 4.5 reviews them: the share of the referrals a Member made
    # that Bellbird accepted within two hours, and the share of its accepted referrals first visited within 24 hours;
    # clause 4.2's response (accepted or declined within two hours), or acceptances over accepted referrals, are other
    # figures for every Member
    s34 = re.sub(r"\s+", " ", pdf_text(os.path.join(target, K.SCHED34)))
    ok("the share of the referrals the Member made in the quarter that Bellbird accepted within two hours of the "
       "Member making them" in s34 and "the share of the Member's accepted referrals that Bellbird first visited "
       "within 24 hours of accepting them" in s34, "B clause 4.5 files both shares' bases")
    allrefs = C.referrals(A)
    for m in K.MEMBERS:
        rr = [x for x in allrefs if x["member"] == m]
        ok(len(rr) == b[m]["n"] and all(x["resp"] is not None for x in rr), "B %s every referral has a response" % m)
        responded = sum(1 for x in rr if (x["resp"] - x["sent"]).total_seconds() <= 7200)
        ok(_r1(100.0 * responded / len(rr)) != _r1(b[m]["pct_acc2"]) and
           _r1(100.0 * b[m]["acc2"] / b[m]["acc"]) != _r1(b[m]["pct_acc2"]),
           "B %s responded within two hours, or over accepted referrals, is another figure" % m)
    # the chart's pay weeks (the finance calendar's, Monday to Sunday): the export holds visits on the quarter's days
    # only (Melbourne dates), so the part weeks commencing 29 June and 28 September hold the quarter's days however a
    # straddling week is read, and each panel's total is the Member's payable visits for the quarter under any weeks
    ok(all(K.Q3S <= v["start"].date() <= K.Q3E for v in A.visits), "chart every export record on a quarter day")
    pay = {m: sum(1 for v in vs if v["member"] == m and v["status"] != "cancel") for m in K.MEMBERS}
    ok(all(sum(wk[(m, w)] for w in C.WEEKS) == pay[m] for m in K.MEMBERS), "chart the 14 pay weeks hold every "
                                                                            "payable visit of the quarter")
    dct = openpyxl.load_workbook(os.path.join(target, K.DICTIONARY), data_only=True).worksheets[0]
    ok(any(r[0] == K.CALENDAR and r[1] == "week_commencing" and "pay week" in (r[2] or "")
           for r in dct.iter_rows(values_only=True)), "chart the dictionary defines the calendar's pay week")
    # payable visits are a count of the visits a fee is payable for (clauses 3.2 to 3.5: completed and attempted);
    # completed visits alone are another count for every Member
    ok("the visit is an attempted visit and its fee is" in s34 and "No fee is payable for a visit cancelled" in s34,
       "chart clauses 3.4 and 3.5 make attempted visits payable and cancelled visits not")
    done = {m: sum(1 for v in vs if v["member"] == m and v["status"] == "done") for m in K.MEMBERS}
    ok(all(done[m] != pay[m] for m in K.MEMBERS), "chart completed visits alone miss every Member's total")
    # the cracker sheet: every handling of the three devices (re-posts, telephoned referrals, intake files), scored
    # against the 35 ask figures (15 amounts, 15 referral figures, 5 chart totals)
    prof = {}
    dd = C.dedupe(A.visits)
    for rp, rg, lk in itertools.product((False, True), repeat=3):
        vv = dd if rp else list(A.visits)
        kept = vv if lk else [v for v in vv if v["client"] != WD.CLIENT_OFFICE]
        ax = C.ask_a(A, visits=kept)
        wx = C.weekly(A, visits=kept)
        bx = C.ask_b(A, refs=C.referrals(A, register=rg), visits=vv, link=lk) if lk else \
            C.ask_b(A, refs=C.referrals(A, register=rg), visits=vv, link=False)
        h = sum(_round(ax[k]) == _round(a[k]) for k in a)
        h += sum(sum(wx[(m, w)] for w in C.WEEKS) == tot[m] for m in K.MEMBERS)
        h += sum((bx[m]["n"] == b[m]["n"]) + (_r1(bx[m]["pct_acc2"]) == _r1(b[m]["pct_acc2"])) +
                 (_r1(bx[m]["pct_v24"]) == _r1(b[m]["pct_v24"])) for m in K.MEMBERS)
        prof[(rp, rg, lk)] = h
    ok(prof[(False, False, False)] <= 2, "cracker: no device handled hits %d of 35" % prof[(False, False, False)])
    ok(prof[(False, False, True)] == 0, "cracker: intake files alone hit nothing")
    ok(prof[(True, False, False)] <= 10, "cracker: re-posts alone hit %d of 35" % prof[(True, False, False)])
    ok(prof[(False, True, False)] <= 12, "cracker: telephoned referrals alone hit %d of 35" % prof[(False, True, False)])
    ok(max(v for k, v in prof.items() if sum(k) == 1) <= 12, "cracker: no single device unlocks more than 12 of 35")
    ok(prof[(True, True, True)] == 35, "cracker: every device handled hits every figure")
    print("cracker profiles (re-posts, telephoned, intake files): %s" % {"".join("RTI"[i] if k[i] else "-" for i in
                                                                         range(3)): v for k, v in prof.items()})
    return dict(a=a, b=b, wk=wk, tot=tot)


# ---------------------------------------------------------------------------------------------------- goldens

def _goldens(golden, B, G, asks):
    ok(sorted(os.listdir(golden)) == sorted([K.GOLD_BOOK, K.GOLD_CHART]), "the two goldens and nothing else")
    wb = openpyxl.load_workbook(os.path.join(golden, K.GOLD_BOOK), data_only=True)
    ok([ws.title for ws in wb.worksheets] == ["Booking", "Owed Jul-Sep 2026", "Referrals Jul-Sep 2026", "Sources"],
       "golden workbook sheets")
    by_name = {K.NAME[m]: m for m in K.MEMBERS}
    ws = wb["Booking"]
    got = {}
    for r in ws.iter_rows(values_only=True):
        if r[0] in by_name:
            got[by_name[r[0]]] = r
    ok(all(got[m][5] == GOLD[m] and got[m][2] == REQ50[m] and got[m][1] == B.plan_q[m] and got[m][3] == B.capacity[m]
           and got[m][4] == GOLD[m] // K.LOT for m in K.MEMBERS), "golden booking sheet carries the call")
    ws = wb["Owed Jul-Sep 2026"]
    for r in ws.iter_rows(values_only=True):
        if r[0] in by_name:
            m = by_name[r[0]]
            ok([r[1], r[2], r[3]] == [_round(asks["a"][(m, mo)]) for mo in (7, 8, 9)], "golden owed row %s" % m)
    ws = wb["Referrals Jul-Sep 2026"]
    for r in ws.iter_rows(values_only=True):
        if r[0] in by_name:
            m = by_name[r[0]]
            x = asks["b"][m]
            ok(list(r[1:7]) == [x["n"], x["acc2"], _r1(x["pct_acc2"]), x["acc"], x["v24"], _r1(x["pct_v24"])],
               "golden referral row %s" % m)
    png = open(os.path.join(golden, K.GOLD_CHART), "rb").read()
    ok(png[:8] == b"\x89PNG\r\n\x1a\n" and b"atplotlib" not in png, "chart PNG without a writer signature")


# ---------------------------------------------------------------------------------------------------- gates

def _gates(target, golden, B, G, A, asks, W, texts, meta_path):
    files = sorted(os.listdir(target))
    ok(len(files) >= 10, "input gate: %d files" % len(files))
    fmts = {f.rsplit(".", 1)[1] for f in files}
    ok(len(fmts) >= 3, "input gate: formats %s" % sorted(fmts))
    rows = {f: sum(1 for _ in open(os.path.join(target, f), encoding="utf-8")) - (1 if f.endswith(".csv") else 0)
            for f in files if f.endswith((".csv", ".jsonl"))}
    ok(max(rows.values()) >= 25000, "input gate: largest file %d rows" % max(rows.values()))
    meta = json.load(open(meta_path))
    ok(sorted(meta["distractor_files"]) == sorted(K.DISTRACTORS) and len(K.DISTRACTORS) >= 2, "distractors declared")
    ok(all(d in files and "distract" not in d for d in K.DISTRACTORS), "distractors shipped, unnamed as such")
    # the call and every ask unchanged with each distractor deleted, and the call with every ask file deleted
    base = (G["booking"], {m: round(G["reqs"][m], 6) for m in K.MEMBERS})
    tmp = tempfile.mkdtemp(prefix="t142chk")
    try:
        for d in list(K.DISTRACTORS) + ["ASKS"]:
            sub = os.path.join(tmp, d.replace(".", "_"))
            os.makedirs(sub)
            drop = set(K.ASK_FILES) | set(K.DISTRACTORS) if d == "ASKS" else {d}
            for f in files:
                if f not in drop:
                    shutil.copy(os.path.join(target, f), os.path.join(sub, f))
            Bx = C.load(sub)
            Gx = C.golden(Bx)
            ok((Gx["booking"], {m: round(Gx["reqs"][m], 6) for m in K.MEMBERS}) == base,
               "the call unchanged without %s" % ("the ask files and distractors" if d == "ASKS" else d))
            if d != "ASKS":
                Ax = C.load_asks(sub)
                ok(C.ask_a(Ax) == asks["a"] and C.ask_b(Ax) == C.ask_b(A), "the asks unchanged without %s" % d)
    finally:
        shutil.rmtree(tmp)
    # zero device rows in the call's files
    for f in K.CALL_FILES:
        t = texts.get(f) if f in texts else open(os.path.join(target, f), encoding="utf-8").read()
        ok(not re.search(r"\bBV\d{7}\b|\bC44\d\d\b|C3920|IL26-|IHR-\d", t), "separation: %s carries no ask row" % f)
    # registers: the manifest names the shipped set; the dictionary names every data file it covers
    man = C._csv(os.path.join(target, K.MANIFEST))
    ok(sorted(r["file"] for r in man) == sorted(f for f in files if f != K.MANIFEST), "H9 manifest equals the set")
    for r in man:
        f = r["file"]
        path = os.path.join(target, f)
        if f.endswith((".csv", ".jsonl")):
            want = _count(rows[f], "row")
        elif f.endswith(".pdf"):
            from pypdf import PdfReader
            want = _count(len(PdfReader(path).pages), "page")
        elif f.endswith(".docx"):
            with zipfile.ZipFile(path) as z:
                want = _count(int(re.search(r"<Pages>(\d+)</Pages>", z.read("docProps/app.xml").decode()).group(1)),
                              "page")
        else:
            want = _xlsx_size(path)
        ok(r["rows_or_pages"] == want, "H9 manifest gives %s for %s (%s counted)" % (r["rows_or_pages"], f, want))
    wb = openpyxl.load_workbook(os.path.join(target, K.DICTIONARY), data_only=True)
    data_ext = (".csv", ".xlsx", ".jsonl")
    dfiles = {r[0] for r in wb.worksheets[0].iter_rows(values_only=True) if r[0] and r[0].endswith(data_ext)}
    ok(set(f for f in files if f.endswith(data_ext) and f not in (K.DICTIONARY, K.SURVEY)) <= dfiles,
       "H9 the dictionary covers every data file: missing %s" %
       sorted(set(f for f in files if f.endswith(data_ext) and f not in (K.DICTIONARY, K.SURVEY)) - dfiles))
    # realism sweep
    adm = collections.Counter((e["member"], e["adm"].year, e["adm"].month) for e in B.eps)
    series = {m: [adm[(m, y, mo)] for (y, mo) in sorted({(e["adm"].year, e["adm"].month) for e in B.eps})]
              for m in K.MEMBERS}
    ok(len({tuple(v) for v in series.values()}) == 5, "realism: Members' monthly admissions differ")
    for m, v in series.items():
        runs = max(len(list(g)) for _, g in itertools.groupby(v))
        ok(runs <= 2, "realism: %s monthly admissions never flat for three months" % m)
    ok(all(v % 50 for v in B.plan_q.values()) and all(v % 50 for v in B.capacity.values()),
       "realism: no round plan or capacity figure")
    notes = collections.Counter(x["note"] for x in W.reg)
    ok(max(notes.values()) <= 12 and len(notes) >= 300, "realism: intake notes vary (max repeat %d, %d distinct)" %
       (max(notes.values()), len(notes)))
    callers_ok = all((x["caller"] in WD.FAMILY) == (x["kind"] in ("Enquiry", "Feedback") and x["org"] in
                                                    [K.NAME[m] for m in K.MEMBERS]) or x["org"] not in
                     [K.NAME[m] for m in K.MEMBERS] for x in W.reg)
    ok(callers_ok, "realism: Members' staff ring about rosters and referrals, families about visits")
    card = json.load(open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(target))),
                                       ".claude", "skills", "fingerprint", "cards", "task142.json")))
    people = set(card["world"]["people"])
    named = set()
    for t in texts.values():
        for p in K.PEOPLE.values():
            if p in t:
                named.add(p)
    ok(named and named <= people, "realism: every named person is on the card %s" % sorted(named - people))
    for d in (target, golden):
        for f in os.listdir(d):
            data = open(os.path.join(d, f), "rb").read()
            ok(not re.search(rb"openpyxl|python-docx|ReportLab|reportlab|atplotlib|Pillow|PIL\b", data),
               "realism: %s carries no writer signature" % f)
    stamp = int(K.PACK_STAMP.timestamp())
    ok(all(int(os.path.getmtime(os.path.join(target, f))) == stamp for f in files), "every pack mtime the export date")
    for f, col, last in ((K.EPISODES, "admission_date", K.D1), (K.CENSUS, "census_date", K.D1),
                         (K.PORTAL, "sent_at", K.Q3E)):
        mx = max(r[col][:10] for r in C._csv(os.path.join(target, f)))
        ok(mx <= last.isoformat(), "dates: %s runs to %s" % (f, mx))
    regmax = max((r["Date"].date() if isinstance(r["Date"], dt.datetime) else r["Date"]) for r in A.register)
    ok(regmax <= dt.date(2026, 10, 30), "dates: the intake log runs to its extract date")
    ok(max(v["rec"] for v in A.visits).date() <= dt.date(2026, 10, 5), "dates: visit records posted by the extract")
