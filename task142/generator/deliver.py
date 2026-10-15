"""The goldens, from the shipped bundle only: hith_overflow_block_2027q1.xlsx (the booking with each Member's
90th-percentile requirement, what the Alliance owes Bellbird for July to September 2026 by Member and month, and
Bellbird's referral response by Member) and bellbird_weekly_visits_2026q3.png (payable visits by pay week and Member);
then generator/figures.json, every figure as submission.md states it."""
from __future__ import annotations

import datetime as dt
import json
import math
import os

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import calc as C
import common as K
import postprocess as PP

MONTHS = (7, 8, 9)
INK = "1F2A30"
GREY = "5B646B"
HEAD_FILL = "DCE6E0"
TOTAL_FILL = "F2F5F3"
THIN = Side(style="thin", color="A9B4AD")


def r1(x):
    return math.floor(x * 10 + 0.5) / 10


def dollars(x):
    return int(math.floor(x + 0.5))


def fmt_int(n):
    return "{:,}".format(int(n))


def compute(target):
    B = C.load(target)
    G = C.golden(B)
    D = C.default_path(B)
    A = C.load_asks(target)
    owed = C.ask_a(A)
    weekly = C.weekly(A)
    resp = C.ask_b(A)
    return B, G, D, A, owed, weekly, resp


# ---------------------------------------------------------------------------------------------------- workbook

def _cell(ws, row, col, value, fmt=None, bold=False, fill=None, align=None, italic=False, color=INK, size=10):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name="Calibri", size=size, bold=bold, italic=italic, color=color)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if align:
        c.alignment = align
    return c


def _header(ws, row, labels, widths):
    for j, (lab, w) in enumerate(zip(labels, widths), 1):
        c = _cell(ws, row, j, lab, bold=True, fill=HEAD_FILL,
                  align=Alignment(wrap_text=True, vertical="bottom", horizontal="left" if j == 1 else "right"))
        c.border = Border(bottom=THIN)
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[row].height = 44


def _span(ws, row, text, last_col, widths, size, **kw):
    """Text across the table's width, wrapped, the row as tall as its lines."""
    c = _cell(ws, row, 1, text, size=size, align=Alignment(wrap_text=True, vertical="top"), **kw)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    chars = sum(widths[:last_col]) * 1.18 * 10 / size
    lines = max(1, math.ceil(len(text) / chars))
    ws.row_dimensions[row].height = lines * size * 1.32 + 2
    return c


def _title(ws, title, sub, last_col, widths):
    _cell(ws, 1, 1, title, bold=True, size=13)
    _span(ws, 2, sub, last_col, widths, 9.5, italic=True, color=GREY)


def _notes(ws, row, lines, last_col, widths):
    for t in lines:
        _span(ws, row, t, last_col, widths, 9, color=GREY)
        row += 1
    return row


def _print_setup(ws, last_col, last_row, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = "A1:%s%d" % (get_column_letter(last_col), last_row)
    ws.oddFooter.left.text = "Ironbark Health Alliance, Hospital in the Home program office"
    ws.oddFooter.right.text = "Page &P of &N"


def write_workbook(path, B, G, owed, resp):
    wb = openpyxl.Workbook()
    # ---- the booking
    ws = wb.active
    ws.title = "Booking"
    last = [m for m in K.MEMBERS if G["booking"][m] > 0][-1]
    wd = [24, 17, 17, 16, 10, 12]
    _title(ws, "Bellbird overflow block, January to March 2027: %s takes the last lots" % K.SHORT[last],
           "Booking under Schedule 2 (Variation 9) of the Bellbird agreement: 850 visits in lots of 25, Members in the "
           "clause 2.2 order. Prepared by the Hospital in the Home program office for lodging by 6 November 2026.", 6, wd)
    hr = 4
    _header(ws, hr, ["Member", "Expected visits, January to March 2027 (operating plan)",
                     "90th-percentile requirement (visits, nearest 50)", "In-house visits (capacity return)",
                     "Lots booked", "Visits booked"], wd)
    row = hr + 1
    for m in K.MEMBERS:
        _cell(ws, row, 1, K.NAME[m])
        _cell(ws, row, 2, B.plan_q[m], "#,##0")
        _cell(ws, row, 3, K.round_to(G["reqs"][m], 50), "#,##0")
        _cell(ws, row, 4, B.capacity[m], "#,##0")
        _cell(ws, row, 5, G["booking"][m] // K.LOT, "0")
        _cell(ws, row, 6, G["booking"][m], "#,##0")
        row += 1
    for j in range(1, 7):
        ws.cell(row=row, column=j).fill = PatternFill("solid", fgColor=TOTAL_FILL)
        ws.cell(row=row, column=j).border = Border(top=THIN)
    _cell(ws, row, 1, "Block", bold=True, fill=TOTAL_FILL)
    _cell(ws, row, 2, sum(B.plan_q.values()), "#,##0", bold=True, fill=TOTAL_FILL)
    _cell(ws, row, 4, sum(B.capacity.values()), "#,##0", bold=True, fill=TOTAL_FILL)
    _cell(ws, row, 5, sum(G["booking"].values()) // K.LOT, "0", bold=True, fill=TOTAL_FILL)
    _cell(ws, row, 6, sum(G["booking"].values()), "#,##0", bold=True, fill=TOTAL_FILL)
    end = row
    first_three = sum(G["booking"][m] for m in K.MEMBERS[:3])
    row = _notes(ws, row + 2, [
        "Requirement: the nurse visits a Member's admitted Hospital in the Home patients receive on 1 January to 31 March "
        "2027, patients admitted before the quarter included for their days in it (clause 2.3); its expected value is the "
        "plan (clause 2.5).",
        "The 90th percentile is built from the plan's admissions by stream from October 2026, each admission given a stay "
        "drawn from the Alliance's completed episodes of that stream (admitted 1 April 2024 to 30 June 2026), visits "
        "counted on the quarter's days by the July 2025 visit schedule.",
        "Lots cover the unrounded percentile's excess over in-house visits (clause 2.6); the requirement is shown to the "
        "nearest 50 visits.",
        "Northgate, Riverbend and Westmere take %s visits in full; %s takes the %d visits left and %s books none "
        "(clause 2.7)." % (fmt_int(first_three), K.SHORT[last], G["booking"][last],
                           ", ".join(K.SHORT[m] for m in K.MEMBERS if G["booking"][m] == 0)),
    ], 6, wd)
    ws.freeze_panes = ws.cell(row=hr + 1, column=2)
    _print_setup(ws, 6, row)
    # ---- what the Alliance owes Bellbird
    ws = wb.create_sheet("Owed Jul-Sep 2026")
    wd = [24, 14, 14, 15, 16]
    _title(ws, "The Alliance owes Bellbird $%s for July to September 2026" % fmt_int(dollars(sum(owed.values()))),
           "Schedule 3 fees from 1 July 2026 (GST-free), by Member and by the month in which each visit started, "
           "Melbourne time. Bellbird visit export of 5 October 2026, each visit counted once.", 5, wd)
    hr = 4
    _header(ws, hr, ["Member", "July 2026", "August 2026", "September 2026", "July to September"], wd)
    row = hr + 1
    money = '"$"#,##0'
    for m in K.MEMBERS:
        _cell(ws, row, 1, K.NAME[m])
        for j, mo in enumerate(MONTHS, 2):
            _cell(ws, row, j, dollars(owed[(m, mo)]), money)
        _cell(ws, row, 5, dollars(sum(owed[(m, mo)] for mo in MONTHS)), money)
        row += 1
    for j in range(1, 6):
        ws.cell(row=row, column=j).fill = PatternFill("solid", fgColor=TOTAL_FILL)
        ws.cell(row=row, column=j).border = Border(top=THIN)
    _cell(ws, row, 1, "All Members", bold=True, fill=TOTAL_FILL)
    for j, mo in enumerate(MONTHS, 2):
        _cell(ws, row, j, dollars(sum(owed[(m, mo)] for m in K.MEMBERS)), money, bold=True, fill=TOTAL_FILL)
    _cell(ws, row, 5, dollars(sum(owed.values())), money, bold=True, fill=TOTAL_FILL)
    row = _notes(ws, row + 2, [
        "Completed visits at the fee for the kind of visit plus the one highest loading (clause 3.3: weekday before 7.00 am "
        "or from 6.00 pm 25 per cent, Saturday 40, Sunday 60, Victorian public holiday 100); attempted visits %s, no "
        "loading (3.4); cancelled visits nil (3.5)." % ("$%.2f" % K.ATTEMPT_FEE),
        "Bellbird's export is stamped in UTC; dates and loadings are on Melbourne time (3.6). Friday 25 September 2026 is "
        "a public holiday in the Alliance finance calendar.",
        "A record re-posted by Bellbird's app under a new visit number (the same nurse and patient, overlapping times) is "
        "the same visit and is counted once. Amounts and totals are each rounded to the dollar from the unrounded fees.",
        "Visits to patients telephoned through during the portal outage of 10 to 13 August were recorded under the "
        "Alliance's billing account (C4466) with Bellbird's intake log number until 14 August; each is charged to the "
        "patient's Member from the intake log.",
    ], 5, wd)
    ws.freeze_panes = ws.cell(row=hr + 1, column=2)
    _print_setup(ws, 5, row)
    # ---- referral response
    ws = wb.create_sheet("Referrals Jul-Sep 2026")
    wd = [24, 12, 13, 13, 12, 15, 17]
    lo2 = min(r1(resp[m]["pct_acc2"]) for m in K.MEMBERS)
    hi2 = max(r1(resp[m]["pct_acc2"]) for m in K.MEMBERS)
    _title(ws, "Bellbird accepted %.1f to %.1f per cent of Members' referrals within two hours, July to September 2026"
           % (lo2, hi2),
           "Referrals each Member made to Bellbird (Schedule 4, clause 4.1): those sent through the Alliance referral "
           "portal and those telephoned to Bellbird's intake line while the portal was down from 10 to 13 August.", 7, wd)
    hr = 4
    _header(ws, hr, ["Member", "Referrals made", "Accepted within 2 hours", "Accepted within 2 hours (%)",
                     "Accepted referrals", "First visited within 24 hours of acceptance",
                     "Accepted referrals first visited within 24 hours (%)"], wd)
    row = hr + 1
    for m in K.MEMBERS:
        x = resp[m]
        _cell(ws, row, 1, K.NAME[m])
        _cell(ws, row, 2, x["n"], "0")
        _cell(ws, row, 3, x["acc2"], "0")
        _cell(ws, row, 4, r1(x["pct_acc2"]), "0.0")
        _cell(ws, row, 5, x["acc"], "0")
        _cell(ws, row, 6, x["v24"], "0")
        _cell(ws, row, 7, r1(x["pct_v24"]), "0.0")
        row += 1
    end = row - 1
    row = _notes(ws, row + 1, [
        "Clocks: from the Member making the referral to Bellbird's acceptance (clause 4.2), and from acceptance to "
        "Bellbird's first attendance at the patient's home, completed or not at home (clause 4.3), all on Melbourne time; "
        "Bellbird's platform times are converted from UTC.",
        "Telephoned referrals are taken from Bellbird's intake log; Bellbird's first visits to those patients are under "
        "the intake log number until the desk linked the files on 14 August.",
        "Shares are of the referrals made and of the accepted referrals respectively (clause 4.5), to one decimal.",
    ], 7, wd)
    ws.freeze_panes = ws.cell(row=hr + 1, column=2)
    _print_setup(ws, 7, row)
    # ---- notes
    ws = wb.create_sheet("Sources")
    wd = [58, 60, 14]
    _title(ws, "Sources and as-of dates", "Files as received by the Hospital in the Home program office.", 3, wd)
    _header(ws, 4, ["File", "Used for", "As of"], wd)
    src = [(K.SCHED2, "The block, the order, the requirement and the lots", "2 October 2026"),
           (K.PLANS, "Expected visits and planned admissions by stream", "14 October 2026"),
           (K.CAPACITY, "In-house visits, January to March 2027", "19 October 2026"),
           (K.EPISODES, "Stays of completed episodes by stream", "30 September 2026"),
           (K.CROSSWALK, "Streams for admissions before 1 July 2025", "1 July 2025"),
           (K.STREAMDOC, "Visit schedule by stream", "16 June 2025"),
           (K.PATHWAY, "Bone and joint infection at Northgate and Westmere from 1 October 2026", "19 August 2026"),
           (K.SCHED34, "Fees, loadings, referral and first-visit service levels", "3 July 2026"),
           (K.BB_VISITS, "Visits, July to September 2026", "5 October 2026"),
           (K.PORTAL, "Portal referrals, July to September 2026", "2 October 2026"),
           (K.REGISTER, "Referrals telephoned during the portal outage", "30 October 2026"),
           (K.ACCOUNTS, "Bellbird client accounts by Member", "5 October 2026"),
           (K.CALENDAR, "Public holidays and pay weeks", "12 June 2026")]
    row = 5
    for f, use, asof in src:
        _cell(ws, row, 1, f)
        _cell(ws, row, 2, use)
        _cell(ws, row, 3, asof, align=Alignment(horizontal="right"))
        row += 1
    _print_setup(ws, 3, row, landscape=False)
    wb.save(path)
    PP.normalize_ooxml(path, K.PEOPLE["requester"], K.PEOPLE["requester"], "2026-11-03T22:41:18Z",
                       "2026-11-04T05:17:42Z", excel=True, company=K.ALLIANCE)


# ---------------------------------------------------------------------------------------------------- chart

PALETTE = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4"]   # validated categorical slots 1 to 5 (light)


def write_chart(path, weekly):
    """Small multiples: one panel per Member in the Schedule's order, each its own colour on a shared scale, and a
    sixth panel with the Members together."""
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib import font_manager
    import matplotlib.pyplot as plt
    for f in ("Carlito-Regular.ttf", "Carlito-Bold.ttf"):
        font_manager.fontManager.addfont("/usr/share/fonts/truetype/crosextra/" + f)
    plt.rcParams.update({"font.family": "Carlito", "font.size": 10, "axes.edgecolor": "#b9beb8",
                         "axes.linewidth": 0.8, "xtick.color": "#52514e", "ytick.color": "#52514e",
                         "axes.labelcolor": "#52514e"})
    weeks = C.WEEKS
    tot = {m: sum(weekly[(m, w)] for w in weeks) for m in K.MEMBERS}
    top = max(K.MEMBERS, key=lambda m: tot[m])
    allv = sum(tot.values())
    allw = [sum(weekly[(m, w)] for m in K.MEMBERS) for w in weeks]
    ymax = max(weekly.values())
    fig, axes = plt.subplots(2, 3, figsize=(11.2, 6.4), dpi=150, sharex=True)
    fig.patch.set_facecolor("#fcfcfb")
    x = list(range(len(weeks)))
    labels = ["%d %s" % (w.day, K.MON3[w.month]) for w in weeks]
    panels = [(m, [weekly[(m, w)] for w in weeks], PALETTE[i], "%s, %d visits" % (K.SHORT[m], tot[m]))
              for i, m in enumerate(K.MEMBERS)] + [(None, allw, "#52514e", "All Members, %s visits" % fmt_int(allv))]
    for ax, (m, y, col, title) in zip(axes.flat, panels):
        ax.set_facecolor("#fcfcfb")
        ax.bar(x, y, width=0.72, color=col, zorder=3)
        k = max(range(len(y)), key=lambda i: (y[i], -i))
        ax.annotate(str(y[k]), xy=(x[k], y[k]), xytext=(0, 3), textcoords="offset points", ha="center", fontsize=8.5,
                    color="#3a3a38")
        ax.set_title(title, loc="left", fontsize=10.5, color="#1f2a30", fontweight="bold" if m == top else "normal")
        lim = (ymax // 5 + 2) * 5 if m else (max(allw) // 10 + 2) * 10
        ax.set_ylim(0, lim)
        ax.yaxis.grid(True, color="#e6e6e1", linewidth=0.7)
        ax.set_axisbelow(True)
        for s in ("top", "right"):
            ax.spines[s].set_visible(False)
        ax.set_xticks(x[::2])
        ax.set_xticklabels(labels[::2], fontsize=8.5)
        ax.tick_params(axis="x", length=2)
    for ax in axes[:, 0]:
        ax.set_ylabel("Payable visits in the week")
    fig.suptitle("%s had the most of Bellbird's %s payable visits from July to September 2026" % (K.SHORT[top],
                                                                                                fmt_int(allv)),
                 x=0.06, ha="left", fontsize=14, fontweight="bold", color="#1f2a30", y=0.975)
    fig.text(0.06, 0.905, "Visits completed or attempted, by pay week commencing Monday (29 June and 28 September are part "
             "weeks). Member panels share one scale; each panel labels its busiest week.", fontsize=9.5,
             color="#52514e")
    fig.text(0.06, 0.018, "Source: Bellbird visit export of 5 October 2026 (UTC, read on Melbourne dates), each visit "
             "counted once; pay weeks from the Alliance finance calendar 2026-27.", fontsize=8.5, color="#7a7a75")
    fig.subplots_adjust(left=0.06, right=0.985, top=0.84, bottom=0.09, hspace=0.42, wspace=0.16)
    fig.savefig(path, facecolor=fig.get_facecolor(), metadata={"Software": None})
    plt.close(fig)
    return tot, top, allv


# ---------------------------------------------------------------------------------------------------- figures

def figures(B, G, D, owed, resp, tot, top, allv):
    S = K.SHORT
    first_three = sum(G["booking"][m] for m in K.MEMBERS[:3])
    last = [m for m in K.MEMBERS if G["booking"][m] > 0][-1]
    f = {
        "call": {"booking": {S[m]: fmt_int(G["booking"][m]) for m in K.MEMBERS},
                 "lots": {S[m]: str(G["booking"][m] // K.LOT) for m in K.MEMBERS},
                 "block": "850", "lot": "25", "last_member": S[last], "left": str(G["left"]),
                 "first_three_visits": fmt_int(first_three)},
        "components": {"requirement_nearest_50": {S[m]: fmt_int(K.round_to(G["reqs"][m], 50)) for m in K.MEMBERS},
                       "lots_needed_first_three": {S[m]: str(G["need"][m]) for m in K.MEMBERS[:3]}},
        "steps": {"plan_q": {S[m]: fmt_int(B.plan_q[m]) for m in K.MEMBERS},
                  "capacity": {S[m]: fmt_int(B.capacity[m]) for m in K.MEMBERS},
                  "pool_first_admission": "1 April 2024", "pool_last_admission": "30 June 2026",
                  "quarter": "1 January to 31 March 2027", "percentile": "90",
                  "schedule2_clauses": ["2.2", "2.3", "2.4", "2.5", "2.6", "2.7"]},
        "default_path": {"booking": {S[m]: fmt_int(D["booking"][m]) for m in K.MEMBERS}, "left": fmt_int(D["left"]),
                         "booked": fmt_int(K.BLOCK - D["left"]),
                         "requirement_nearest_50": {S[m]: fmt_int(K.round_to(D["reqs"][m], 50)) for m in K.MEMBERS}},
        "deliverables": {
            "owed": {S[m]: {K.MONTH_NAME[mo]: "$" + fmt_int(dollars(owed[(m, mo)])) for mo in MONTHS} for m in K.MEMBERS},
            "owed_member_quarter": {S[m]: "$" + fmt_int(dollars(sum(owed[(m, mo)] for mo in MONTHS))) for m in K.MEMBERS},
            "owed_month_total": {K.MONTH_NAME[mo]: "$" + fmt_int(dollars(sum(owed[(m, mo)] for m in K.MEMBERS)))
                                 for mo in MONTHS},
            "owed_total": "$" + fmt_int(dollars(sum(owed.values()))),
            "referrals": {S[m]: {"made": str(resp[m]["n"]), "accepted_2h": str(resp[m]["acc2"]),
                                 "accepted_2h_pct": "%.1f" % r1(resp[m]["pct_acc2"]), "accepted": str(resp[m]["acc"]),
                                 "visited_24h": str(resp[m]["v24"]), "visited_24h_pct": "%.1f" % r1(resp[m]["pct_v24"])}
                          for m in K.MEMBERS},
            "chart_totals": {S[m]: str(tot[m]) for m in K.MEMBERS}, "chart_top": S[top], "chart_all": fmt_int(allv),
            "chart_weeks": "14", "outage": "10 to 13 August 2026", "linked": "14 August 2026"},
    }
    return f


def main(target, golden, figures_path):
    os.makedirs(golden, exist_ok=True)
    B, G, D, A, owed, weekly, resp = compute(target)
    write_workbook(os.path.join(golden, K.GOLD_BOOK), B, G, owed, resp)
    tot, top, allv = write_chart(os.path.join(golden, K.GOLD_CHART), weekly)
    stamp = dt.datetime(2026, 11, 4, 5, 17, 42, tzinfo=dt.timezone.utc).timestamp()
    for fn in os.listdir(golden):
        os.utime(os.path.join(golden, fn), (stamp, stamp))
    fj = figures(B, G, D, owed, resp, tot, top, allv)
    with open(figures_path, "w", encoding="utf-8") as fh:
        json.dump(fj, fh, indent=2, ensure_ascii=False)
        fh.write("\n")
    return fj


if __name__ == "__main__":
    import sys
    main(sys.argv[1], sys.argv[2], sys.argv[3])
