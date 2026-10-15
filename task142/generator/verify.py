"""Independent verifier for task142. Reads only the shipped bundle (target/ and the goldens) with its own parsing and
its own methods, imports nothing from the generator, and recomputes every graded figure:

* the call: each Member's January to March 2027 requirement from the operating plan's admissions and the stays of the
  Alliance's completed episodes, its distribution computed exactly by Panjer's recursion (the generator uses the
  Fourier transform), the 90th percentile, the lots of Schedule 2 and the booking in its order;
* the default path (plan plus a Poisson margin) and its booking;
* the asks: Schedule 3 amounts by Member and month (clocks through zoneinfo, a re-posted record found by the overlap of
  one nurse's records for one patient), payable visits by pay week, and Schedule 4's referral response (portal and
  intake log, the intake desk's files linked by log number).

Constants of the agreement (block, lot, order, fees, loadings) are read from the PDFs' text."""
from __future__ import annotations

import csv
import datetime as dt
import json
import math
import os
import re
import sys
from zoneinfo import ZoneInfo

import numpy as np
import openpyxl
from pypdf import PdfReader

MEL = ZoneInfo("Australia/Melbourne")
UTC = dt.timezone.utc


def _pdf(path):
    return re.sub(r"\s+", " ", " ".join(p.extract_text() or "" for p in PdfReader(path).pages))


def _rows(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def _find(folder, prefix):
    hits = [f for f in os.listdir(folder) if f.startswith(prefix)]
    assert len(hits) == 1, (prefix, hits)
    return os.path.join(folder, hits[0])


def _date(s):
    return dt.date(int(s[:4]), int(s[5:7]), int(s[8:10]))


def _visits_by_day(schedule, n):
    """Visits on stay days 0..n-1 under the visit schedule's wording."""
    out = np.ones(n)
    if "every second day" in schedule:
        out[1::2] = 0
    m = re.search(r"Two visits on each of the first (\w+) days", schedule)
    if m:
        k = {"two": 2, "three": 3}[m.group(1)]
        out[:k] = 2
    return out


def call(target):
    s2 = _pdf(_find(target, "bellbird_agreement_schedule2"))
    block = int(re.search(r"will hold (\d+) nurse visits", s2).group(1))
    lot = int(re.search(r"the lots of (\d+) visits", s2).group(1))
    order = [x.strip() for x in re.search(r"in this order: ([^.]+)\.", s2).group(1).split(",")]
    sd = _pdf(_find(target, "hith_program_streams_and_visit_schedule"))
    sched = {}
    for st in ["H0%d" % i for i in range(1, 9)]:
        m = re.search(r"\b%s\b (.*?)(?= H0\d\b| Days at home)" % st, sd)
        sched[st] = m.group(1) if m else ""
    xw = {(r["old_program"], r["principal_diagnosis_category"]): r["new_stream"]
          for r in _rows(_find(target, "hith_program_stream_crosswalk"))}
    pool = {}
    open_n = 0
    for r in _rows(_find(target, "hith_episodes_")):
        code = r["program_stream"]
        s = code if re.fullmatch(r"H0\d", code) else (xw.get((code, r["principal_diagnosis"][:3])) or xw[(code, "")])
        if not r["separation_date"]:
            open_n += 1
            continue
        a = _date(r["admission_date"])
        if dt.date(2024, 4, 1) <= a <= dt.date(2026, 6, 30):
            pool.setdefault(s, []).append((_date(r["separation_date"]) - a).days + 1)
    # the plans: admissions by stream and month, visits for January to March 2027
    wb = openpyxl.load_workbook(_find(target, "hith_operating_plans"), data_only=True)
    adm, plan = {}, {}
    names = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        member = rows[0][0].split(":")[0]
        names[ws.title] = member
        hi = next(i for i, r in enumerate(rows) if r[0] == "Stream")
        months = []
        for h in rows[hi][3:15]:
            mon, yy = h.split()[:2]
            months.append(dt.datetime.strptime("1 %s 20%s" % (mon, yy), "%d %b %Y").date())
        q = 0
        for r in rows[hi + 1:]:
            if not r[0] or not str(r[0]).startswith("H"):
                continue
            for mo, v in zip(months, r[3:15]):
                if r[2] == "Admissions":
                    adm.setdefault(member, {}).setdefault(r[0], {})[mo] = v
                elif mo.year == 2027 and mo.month <= 3:
                    q += v
        plan[member] = q
    wb = openpyxl.load_workbook(_find(target, "hith_inhouse_capacity_return"), data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r[0] == "Member")
    col = next(j for j, h in enumerate(rows[hi]) if h and str(h).startswith("In-house visits"))
    cap = {r[0]: r[col] for r in rows[hi + 1:] if r[0] in order}
    # the quarter: 1 January to 31 March 2027; admissions from 1 October 2026 (longer stays end before it otherwise)
    q0, q1 = dt.date(2027, 1, 1), dt.date(2027, 3, 31)
    W = (q1 - q0).days + 1
    req, p90s = {}, {}
    for member in order:
        sev = np.zeros(400)
        for s, by_month in adm[member].items():
            if s not in pool:
                continue
            L = np.array(pool[s])
            vis = _visits_by_day(sched[s], int(L.max()) + 1)
            cum = np.concatenate([[0.0], np.cumsum(vis)])
            d = dt.date(2026, 10, 1)
            while d <= q1:
                first = d.replace(day=1)
                nxt = (first + dt.timedelta(days=32)).replace(day=1)
                rate = by_month.get(first, 0) / (nxt - first).days
                if rate > 0:
                    o = (d - q0).days
                    lo = np.clip(-o, 0, L)
                    hi_ = np.clip(W - o, 0, L)
                    v = np.where(hi_ > lo, cum[hi_] - cum[lo], 0).astype(int)
                    np.add.at(sev, v, rate / len(L))
                d += dt.timedelta(days=1)
        lam = sev.sum()
        f = sev / lam
        kmax = int(np.nonzero(f)[0].max())
        nmax = 6000
        p = np.zeros(nmax)
        p[0] = math.exp(lam * (f[0] - 1))
        kf = np.arange(kmax + 1) * f[:kmax + 1] * lam
        for n in range(1, nmax):
            k = np.arange(1, min(n, kmax) + 1)
            p[n] = float(np.dot(kf[k], p[n - k])) / n
        assert abs(p.sum() - 1) < 1e-9, (member, p.sum())
        mean = float((np.arange(nmax) * p).sum())
        p90 = int(np.searchsorted(np.cumsum(p), 0.9 - 1e-12))
        p90s[member] = p90
        req[member] = p90 + (plan[member] - mean)
    left = block // lot
    booking = {}
    for member in order:
        excess = req[member] - cap[member]
        n = 0 if excess <= 0 else int(math.ceil(excess / lot - 1e-12))
        take = min(n, left)
        booking[member] = take * lot
        left -= take
    # the default path: plan plus the Poisson 90th percentile
    from scipy.stats import poisson
    dleft = block // lot
    dbook = {}
    for member in order:
        excess = poisson.ppf(0.9, plan[member]) - cap[member]
        n = 0 if excess <= 0 else int(math.ceil(excess / lot - 1e-12))
        dbook[member] = min(n, dleft) * lot
        dleft -= min(n, dleft)
    return dict(order=order, booking=booking, req=req, left=left * lot, default=dbook, default_left=dleft * lot,
                plan=plan, cap=cap, open_episodes=open_n)


def asks(target, keep="first_recorded"):
    """Schedule 3 amounts, payable visits by pay week and the referral response, a re-post kept as `keep` says:
    the first or last recorded, or the earliest or latest start, of a nurse's overlapping records for a patient."""
    s3 = _pdf(_find(target, "bellbird_agreement_schedules3-4"))
    fee = {}
    for code in ("IVT", "WND", "CLN"):
        fee[code] = float(re.search(r"\b%s \$(\d+\.\d\d)" % code, s3).group(1))
    attempt = float(re.search(r"attempted visit and its fee is \$(\d+\.\d\d)", s3).group(1))
    ld = re.search(r"before 7\.00 am or at or after 6\.00 pm, (\d+) per cent; Saturday, (\d+) per cent; Sunday, (\d+) "
                   r"per cent; a Victorian public holiday, (\d+) per cent", s3)
    after, sat, sun, hol = (int(x) / 100 for x in ld.groups())
    acc = {r["client_code"]: r["alliance_member_code"] for r in _rows(_find(target, "bellbird_client_accounts"))}
    names = {r["alliance_member_code"]: r["account_name"] for r in _rows(_find(target, "bellbird_client_accounts"))}
    cal = _rows(_find(target, "alliance_finance_calendar"))
    holidays = {_date(r["date"]) for r in cal if r["public_holiday_vic"]}
    week_of = {_date(r["date"]): _date(r["week_commencing"]) for r in cal}

    org_member = {}
    for code, nm in names.items():
        if code:
            org_member[nm.split(" - ")[0]] = code
    org_member["Westmere Health (Altona Park)"] = "WMH"
    wb = openpyxl.load_workbook(_find(target, "bellbird_intake_register"), data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r[0] == "Log no.")
    head = rows[hi]
    reg = [dict(zip(head, r)) for r in rows[hi + 1:] if r[0]]
    digits = lambda s: re.sub(r"\D", "", str(s)).lstrip("0")
    log_patient = {r["Log no."]: (org_member[r["Organisation"]], digits(r["Patient"])) for r in reg
                   if r["Call type"] == "Referral" and r["Organisation"] in org_member}

    def mel(s):
        return dt.datetime.strptime(s, "%Y-%m-%dT%H:%M:%S").replace(tzinfo=UTC).astimezone(MEL).replace(tzinfo=None)
    status = {"COMPLETE": "C", "Completed": "C", "NOTHOME": "A", "Attempted": "A", "CANCELLED": "X", "Cancelled": "X"}
    recs = []
    with open(_find(target, "bellbird_visits_"), encoding="utf-8") as fh:
        visit_rows = [json.loads(line) for line in fh if line.strip()]
    for r in visit_rows:
        member = acc[r["client_code"]] or log_patient[r["patient_ref"].split("/", 1)[1]][0]
        recs.append(dict(member=member, pref=r["patient_ref"], nurse=r["nurse_id"], type=r["visit_type"],
                         start=mel(r["start_time"]), end=mel(r["end_time"]) if r["end_time"] else None,
                         st=status[r["status"]], rec=r["recorded_at"], no=r["visit_no"]))
    by_rec = lambda x: (x["rec"], x["no"])
    by_start = lambda x: (x["start"], x["rec"], x["no"])
    order, rev = {"first_recorded": (by_rec, False), "last_recorded": (by_rec, True),
                  "earliest_start": (by_start, False), "latest_start": (by_start, True)}[keep]
    recs.sort(key=order, reverse=rev)
    kept, by = [], {}
    for x in recs:
        if x["st"] == "X":
            kept.append(x)
            continue
        lst = by.setdefault((x["nurse"], x["pref"]), [])
        if any(x["start"] < y["end"] and y["start"] < x["end"] for y in lst):
            continue
        lst.append(x)
        kept.append(x)
    q3s, q3e = dt.date(2026, 7, 1), dt.date(2026, 9, 30)
    owed, wk = {}, {}
    for x in kept:
        d = x["start"].date()
        if not (q3s <= d <= q3e) or x["st"] == "X":
            continue
        if x["st"] == "A":
            amt = attempt
        else:
            wd = d.weekday()
            load = hol if d in holidays else sun if wd == 6 else sat if wd == 5 else \
                after if (x["start"].hour < 7 or x["start"].hour >= 18) else 0.0
            amt = round(fee[x["type"]] * (1 + load), 2)
        owed[(x["member"], d.month)] = owed.get((x["member"], d.month), 0.0) + amt
        wk[(x["member"], week_of[d])] = wk.get((x["member"], week_of[d]), 0) + 1
    # referrals: the portal's and those telephoned to the intake line
    refs = []
    for r in _rows(_find(target, "referral_portal_bellbird")):
        sent = dt.datetime.strptime(r["sent_at"], "%Y-%m-%d %H:%M")
        if q3s <= sent.date() <= q3e:
            resp = dt.datetime.strptime(r["bellbird_response_at"], "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=UTC) \
                .astimezone(MEL).replace(tzinfo=None)
            refs.append((r["member_code"], sent, r["status"] == "Accepted", resp, digits(r["patient_no"])))
    for r in reg:
        if r["Call type"] != "Referral" or r["Organisation"] not in org_member:
            continue
        d = r["Date"].date() if isinstance(r["Date"], dt.datetime) else r["Date"]
        if q3s <= d <= q3e:
            sent = dt.datetime.combine(d, dt.datetime.strptime(r["Time"], "%H:%M").time())
            resp = dt.datetime.combine(d, dt.datetime.strptime(r["Outcome time"], "%H:%M").time())
            if resp < sent:
                resp += dt.timedelta(days=1)
            refs.append((org_member[r["Organisation"]], sent, r["Outcome"] == "Accepted", resp, digits(r["Patient"])))
    first = {}
    for x in kept:
        if x["st"] == "X":
            continue
        tail = x["pref"].split("/", 1)[1]
        key = log_patient[tail] if tail in log_patient else (x["member"], digits(tail))
        first.setdefault(key, []).append(x["start"])
    resp_out = {}
    for m in sorted({r[0] for r in refs}):
        rr = [r for r in refs if r[0] == m]
        n = len(rr)
        a2 = sum(1 for r in rr if r[2] and (r[3] - r[1]).total_seconds() <= 7200)
        accepted = [r for r in rr if r[2]]
        v24 = 0
        for r in accepted:
            ts = sorted(t for t in first.get((m, r[4]), []) if t >= r[3])
            if ts and (ts[0] - r[3]).total_seconds() <= 24 * 3600:
                v24 += 1
        resp_out[m] = (n, a2, len(accepted), v24)
    return owed, wk, resp_out


def main(target, golden):
    c = call(target)
    owed, wk, resp = asks(target)
    cents = lambda o: {k: round(v, 2) for k, v in o.items()}
    for keep in ("last_recorded", "earliest_start", "latest_start"):
        o2, w2, r2 = asks(target, keep)
        assert cents(o2) == cents(owed) and w2 == wk and r2 == resp, ("the re-post kept moves a figure", keep)
    wb = openpyxl.load_workbook(_find(golden, "hith_overflow_block"), data_only=True)
    short = {}
    ws = wb["Booking"]
    for r in ws.iter_rows(values_only=True):
        if r[0] in c["order"]:
            assert r[5] == c["booking"][r[0]], ("booking", r[0], r[5], c["booking"][r[0]])
            want = int(math.floor(c["req"][r[0]] / 50 + 0.5)) * 50
            assert r[2] == want, ("requirement", r[0], r[2], want)
            assert r[1] == c["plan"][r[0]] and r[3] == c["cap"][r[0]], ("inputs", r[0])
    assert c["left"] == 0 and sum(c["booking"].values()) == 850, c
    assert list(c["default"].values()) == [100, 50, 50, 25, 25] and c["default_left"] == 600, c["default"]
    code = {"Northgate Health": "NGH", "Riverbend Health": "RBH", "Westmere Health": "WMH", "Kingsdale Health": "KDH",
            "Hillcrest Health": "HCH"}
    ws = wb["Owed Jul-Sep 2026"]
    n_owed = 0
    for r in ws.iter_rows(values_only=True):
        if r[0] in code:
            m = code[r[0]]
            for j, mo in enumerate((7, 8, 9), 1):
                assert r[j] == int(math.floor(owed[(m, mo)] + 0.5)), ("owed", m, mo, r[j], owed[(m, mo)])
                n_owed += 1
    ws = wb["Referrals Jul-Sep 2026"]
    n_ref = 0
    for r in ws.iter_rows(values_only=True):
        if r[0] in code:
            n, a2, acc, v24 = resp[code[r[0]]]
            r1 = lambda x: math.floor(x * 10 + 0.5) / 10
            assert list(r[1:7]) == [n, a2, r1(100 * a2 / n), acc, v24, r1(100 * v24 / acc)], ("referrals", r, n, a2,
                                                                                             acc, v24)
            n_ref += 1
    totals = {}
    for (m, w), v in wk.items():
        totals[m] = totals.get(m, 0) + v
    print("verify: call %s (requirements %s), default %s with %d left; %d owed cells, %d referral rows and %d chart "
          "totals recomputed independently, the same whichever record of a re-post is kept" % (
              "/".join(str(c["booking"][m]) for m in c["order"]),
              "/".join("%.1f" % c["req"][m] for m in c["order"]),
              "/".join(str(v) for v in c["default"].values()), c["default_left"], n_owed, n_ref, len(totals)))
    return c, owed, wk, resp, totals


if __name__ == "__main__":
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "target"),
         sys.argv[2] if len(sys.argv) > 2 else os.path.join(here, "golden"))
