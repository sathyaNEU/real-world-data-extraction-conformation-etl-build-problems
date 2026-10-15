"""Every graded figure from the shipped bundle: each member's 90th-percentile requirement for January to March 2027
and the booking (Schedule 2), every rung and rival of the ladder, and the asks (what the alliance owes Bellbird for
July to September 2026 by member and month, Bellbird's payable visits by week, and Bellbird's referral response by
member). Reads only files under target/ (the constants of Schedules 2 to 4 are checked against the PDFs' text by
verify.py)."""
from __future__ import annotations

import csv
import datetime as dt
import math
import os
import re

import numpy as np
import openpyxl

import common as K
import model as MD


class Bundle:
    pass


def _csv(path):
    with open(path, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def _date(s):
    return dt.date.fromisoformat(s) if s else None


# ---------------------------------------------------------------------------------------------------- loading

def load(target):
    B = Bundle()
    p = lambda f: os.path.join(target, f)
    xw = {}
    for r in _csv(p(K.CROSSWALK)):
        xw[(r["old_program"], r["principal_diagnosis_category"])] = r["new_stream"]
    B.crosswalk = xw
    eps = []
    for r in _csv(p(K.EPISODES)):
        code = r["program_stream"]
        if code in K.STREAMS:
            s = code
        else:
            s = xw.get((code, r["principal_diagnosis"][:3])) or xw[(code, "")]
        eps.append(dict(member=r["member_code"], patient=r["patient_no"], episode=r["episode_no"],
                        adm=_date(r["admission_date"]), sep=_date(r["separation_date"]), mode=r["separation_mode"],
                        stream=s, code=code, dx=r["principal_diagnosis"], campus=r["campus_code"]))
    B.eps = eps
    B.census = _csv(p(K.CENSUS))
    B.adm, B.vis, B.plan_q = load_plans(p(K.PLANS))
    B.capacity = load_capacity(p(K.CAPACITY))
    return B


def load_plans(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    adm = {m: {s: {} for s in K.STREAMS} for m in K.MEMBERS}
    vis = {m: {s: {} for s in K.STREAMS} for m in K.MEMBERS}
    by_short = {v: k for k, v in K.SHORT.items()}
    for ws in wb.worksheets:
        m = by_short[ws.title]
        rows = list(ws.iter_rows(values_only=True))
        hi = next(i for i, r in enumerate(rows) if r[0] == "Stream")
        months = []
        for h in rows[hi][3:15]:
            mon, yy = h.split()[:2]
            mi = next(k for k, v in K.MON3.items() if v == mon)
            months.append(dt.date(2000 + int(yy), mi, 1))
        for r in rows[hi + 1:]:
            if not r[0] or r[0] == "All" or not str(r[0]).startswith("H"):
                continue
            tgt = adm if r[2] == "Admissions" else vis
            for mo, v in zip(months, r[3:15]):
                tgt[m][r[0]][mo] = int(v)
    plan_q = {m: sum(vis[m][s].get(dt.date(2027, mo, 1), 0) for s in K.STREAMS for mo in (1, 2, 3)) for m in K.MEMBERS}
    return adm, vis, plan_q


def load_capacity(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r[0] == "Member")
    col = next(j for j, h in enumerate(rows[hi]) if h and str(h).startswith("In-house visits"))
    by_name = {v: k for k, v in K.NAME.items()}
    out = {}
    for r in rows[hi + 1:]:
        if r[0] in by_name:
            out[by_name[r[0]]] = int(r[col])
    return out


# ---------------------------------------------------------------------------------------------------- stays

def stays(B, first=K.D0, last=K.POOL_LAST_ADMISSION, member=None, include_open=False, include_late=False,
          merge=False, km=False):
    """Home days of the alliance's episodes by stream. The golden pool: completed episodes admitted from 1 April 2024
    to 30 June 2026, every member pooled (all complete by 30 September 2026)."""
    eps = B.eps
    if merge:
        eps = merged(B)
    out = {s: [] for s in K.STREAMS}
    cens = {s: [] for s in K.STREAMS}
    for e in eps:
        if member and e["member"] != member:
            continue
        if e["adm"] < first:
            continue
        if e["sep"] is None:
            if include_open:
                out[e["stream"]].append((K.D1 - e["adm"]).days + 1)
            if km:
                cens[e["stream"]].append((K.D1 - e["adm"]).days + 1)
            continue
        if e["adm"] > last and not (include_late or km):
            continue
        out[e["stream"]].append((e["sep"] - e["adm"]).days + 1)
    res = {s: np.array(v, dtype=int) for s, v in out.items()}
    if km:
        res = {s: km_complete(res[s], np.array(cens[s], dtype=int)) for s in K.STREAMS}
    return res


def km_complete(done, censored):
    """Kaplan-Meier on completed and still-open stays, returned as an integer sample with the estimated
    distribution (10,000 values by quantile)."""
    if len(done) == 0:
        return done
    t_all = np.concatenate([done, censored])
    ev = np.concatenate([np.ones(len(done)), np.zeros(len(censored))])
    ts = np.unique(done)
    S = 1.0
    surv = []
    for t in ts:
        at_risk = (t_all >= t).sum()
        d = ((t_all == t) & (ev == 1)).sum()
        S *= 1 - d / at_risk
        surv.append(S)
    cdf = 1 - np.array(surv)
    cdf = cdf / cdf[-1]
    q = (np.arange(10000) + 0.5) / 10000
    return ts[np.searchsorted(cdf, q)]


def merged(B):
    """Episodes with an interrupted course (a mode 2 separation followed within three days by a new episode of the
    same patient and stream) joined into one stay."""
    by = {}
    for e in sorted(B.eps, key=lambda e: e["adm"]):
        by.setdefault((e["member"], e["patient"], e["stream"]), []).append(e)
    out = []
    for lst in by.values():
        i = 0
        while i < len(lst):
            e = dict(lst[i])
            days_ = (e["sep"] - e["adm"]).days + 1 if e["sep"] else None
            while (i + 1 < len(lst) and lst[i]["mode"] == "2" and lst[i]["sep"] is not None
                   and 2 <= (lst[i + 1]["adm"] - lst[i]["sep"]).days <= 4):
                i += 1
                nxt = lst[i]
                if nxt["sep"] is None:
                    days_ = None
                    e["sep"] = None
                    break
                days_ += (nxt["sep"] - nxt["adm"]).days + 1
                e["sep"] = e["adm"] + dt.timedelta(days=days_ - 1)
            out.append(e)
            i += 1
    return out


def continuation_share(B):
    """Per stream, the share of episodes that resume an interrupted course."""
    by = {}
    for e in sorted(B.eps, key=lambda e: e["adm"]):
        by.setdefault((e["member"], e["patient"], e["stream"]), []).append(e)
    n = {s: 0 for s in K.STREAMS}
    c = {s: 0 for s in K.STREAMS}
    for lst in by.values():
        for i, e in enumerate(lst):
            n[e["stream"]] += 1
            if i > 0 and lst[i - 1]["mode"] == "2" and lst[i - 1]["sep"] is not None and \
                    2 <= (e["adm"] - lst[i - 1]["sep"]).days <= 4:
                c[e["stream"]] += 1
    return {s: (c[s] / n[s] if n[s] else 0.0) for s in K.STREAMS}


# ---------------------------------------------------------------------------------------------------- the call

def requirement(B, m, pool=None, start=K.PATH, recentre=True, adm=None):
    pool = pool if pool is not None else stays(B)
    return MD.requirement(adm or B.adm[m], pool, B.plan_q[m], start=start, recentre=recentre)


def golden(B):
    pool = stays(B)
    reqs, det = {}, {}
    for m in K.MEMBERS:
        req, p90, mean, sd = requirement(B, m, pool)
        reqs[m] = req
        det[m] = dict(req=req, p90=p90, mean=mean, sd=sd, plan=B.plan_q[m], cap=B.capacity[m])
    booking, need, left = MD.book(reqs, B.capacity)
    return dict(reqs=reqs, det=det, booking=booking, need=need, left=left, pool=pool)


def poisson_p90(mu):
    from scipy import stats
    return float(stats.poisson.ppf(0.9, mu))


def rung(B, reqs):
    booking, need, left = MD.book(reqs, B.capacity)
    return dict(reqs=reqs, booking=booking, need=need, left=left)


def default_path(B):
    return rung(B, {m: poisson_p90(B.plan_q[m]) for m in K.MEMBERS})


def history_daily(B):
    """Each member's daily visits from the episode extract (every day from 1 July 2024 to 30 September 2026), by the
    visit schedule; and from the census (1 July 2025 to 30 September 2026)."""
    import pandas as pd
    d0 = dt.date(2024, 7, 1)
    dd = K.days(d0, K.D1)
    idx = {d: i for i, d in enumerate(dd)}
    H = {m: np.zeros(len(dd)) for m in K.MEMBERS}
    for e in B.eps:
        last = e["sep"] or K.D1
        for k in range((last - e["adm"]).days + 1):
            d = e["adm"] + dt.timedelta(days=k)
            if d in idx:
                H[e["member"]][idx[d]] += K.visits_on_day(e["stream"], k)
    ep = pd.DataFrame(H, index=pd.to_datetime(dd))
    C = {}
    for r in B.census:
        key = (r["census_date"], r["member_code"])
        C[key] = C.get(key, 0) + int(r["nurse_visits"])
    cd = K.days(K.CENSUS0, K.D1)
    cen = pd.DataFrame({m: [C.get((d.isoformat(), m), 0) for d in cd] for m in K.MEMBERS}, index=pd.to_datetime(cd))
    return ep, cen


def rivals(B, G=None):
    """Every rung below the decisive one and every rival construction, as each member's 90th-percentile requirement;
    the booking follows from Schedule 2's arithmetic."""
    import pandas as pd
    G = G or golden(B)
    ep, cen = history_daily(B)
    z = 1.2815516
    out = {}
    out["R0_poisson"] = {m: poisson_p90(B.plan_q[m]) for m in K.MEMBERS}
    out["R0_poisson_normal"] = {m: B.plan_q[m] + z * math.sqrt(B.plan_q[m]) for m in K.MEMBERS}
    out["R0_daily_independent"] = {}
    for lab, h in (("ep", ep), ("cen", cen)):
        for k in ("R1_monthly_", "R1_weekly_", "R2_quarterly_", "R1_nb_monthly_"):
            out[k + lab] = {}
    out["R1_monthly_deseasonalised_ep"] = {}
    for m in K.MEMBERS:
        plan = B.plan_q[m]
        d = cen[m]
        out["R0_daily_independent"][m] = plan + z * math.sqrt(90) * d.std() * (plan / 90) / d.mean()
        for lab, h in (("ep", ep[m]), ("cen", cen[m])):
            mo = h.resample("MS").sum()
            qt = h.resample("QS").sum()
            wk = h.resample("W-SUN").sum().iloc[1:-1]
            out["R1_monthly_" + lab][m] = plan + z * math.sqrt(3) * (mo.std() / mo.mean()) * plan / 3
            out["R1_weekly_" + lab][m] = plan + z * math.sqrt(90 / 7) * (wk.std() / wk.mean()) * plan / (90 / 7)
            out["R2_quarterly_" + lab][m] = plan + z * (qt.std() / qt.mean()) * plan
            mu, v = mo.mean(), mo.var()
            kk = mu ** 2 / max(v - mu, 1e-9)
            out["R1_nb_monthly_" + lab][m] = plan + z * math.sqrt(3 * ((plan / 3) + (plan / 3) ** 2 / kk))
        mo = ep[m].resample("MS").sum()
        res = mo - mo.groupby(mo.index.month).transform("mean")
        out["R1_monthly_deseasonalised_ep"][m] = plan + z * math.sqrt(3) * res.std(ddof=1) * (plan / 3) / mo.mean()
    # R2b: each member's own past admissions and own stays (no plan mix), recentred on the plan
    out["R2b_own_history"] = {}
    for m in K.MEMBERS:
        own = stays(B, member=m)
        a0, a1 = dt.date(2025, 7, 1), dt.date(2026, 6, 30)
        n = {s: sum(1 for e in B.eps if e["member"] == m and e["stream"] == s and a0 <= e["adm"] <= a1)
             for s in K.STREAMS}
        adm = {s: {mo: n[s] / 365.0 * K.month_end(mo).day for mo in K.MONTHS_FY} for s in K.STREAMS}
        out["R2b_own_history"][m] = MD.requirement(adm, own, B.plan_q[m])[0]
    # the plan's admissions, stays drawn, but no arrival variation (admissions fixed at the plan)
    out["R3x_fixed_admissions"] = {}
    out["R3x_cohort"] = {}
    out["R3x_no_carry_in"] = {}
    for m in K.MEMBERS:
        mix, mean = MD.quarter_mix(B.adm[m], G["pool"])
        # fixed admissions: variance of the sum of independent stays at the planned numbers (normal approximation)
        var_fixed = 0.0
        W = (K.Q1 - K.Q0).days
        for s in K.STREAMS:
            pool = G["pool"][s]
            if len(pool) == 0:
                continue
            rates = MD.daily_rates(B.adm[m][s], K.PATH, K.Q1)
            offs = np.array([(d - K.Q0).days for d in K.days(K.PATH, K.Q1)])
            for lam, o in zip(rates, offs):
                if lam <= 0:
                    continue
                v = MD.window_visits(s, np.full(len(pool), o), pool, W)
                var_fixed += lam * v.var()
        out["R3x_fixed_admissions"][m] = B.plan_q[m] + z * math.sqrt(var_fixed)
        # cohort: the quarter's admissions with their whole stays
        cm = np.zeros(400)
        for s in K.STREAMS:
            pool = G["pool"][s]
            if len(pool) == 0:
                continue
            rates = MD.daily_rates(B.adm[m][s], K.Q0, K.Q1)
            vals, cnt = np.unique(pool, return_counts=True)
            tot = np.array([sum(K.visits_on_day(s, k) for k in range(L)) for L in vals])
            for lam in rates:
                np.add.at(cm, tot, lam * cnt / cnt.sum())
        pmf = MD.compound(cm)
        x = np.arange(len(pmf))
        mu = (x * pmf).sum()
        out["R3x_cohort"][m] = B.plan_q[m] + (MD.quantile(pmf) - mu)
        # admissions from 1 January only (no patients carried into the quarter)
        mix2, mean2 = MD.quarter_mix(B.adm[m], G["pool"], start=K.Q0)
        pmf2 = MD.compound(mix2)
        x2 = np.arange(len(pmf2))
        out["R3x_no_carry_in"][m] = B.plan_q[m] + (MD.quantile(pmf2) - (x2 * pmf2).sum())
    # each admission given its stream's mean stay (stays not varied), arrivals Poisson, normal approximation
    out["R3x_fixed_stays"] = {}
    # the quarter's admissions with whole stays, compound Poisson variance, normal approximation
    out["R3x_whole_stay_normal"] = {}
    W = (K.Q1 - K.Q0).days
    offs_all = np.array([(d - K.Q0).days for d in K.days(K.PATH, K.Q1)])
    for m in K.MEMBERS:
        var_f, var_w = 0.0, 0.0
        for s in K.STREAMS:
            pool = G["pool"][s]
            if len(pool) == 0:
                continue
            Lbar = int(round(pool.mean()))
            rates = MD.daily_rates(B.adm[m][s], K.PATH, K.Q1)
            for lam, o in zip(rates, offs_all):
                if lam > 0:
                    v = int(MD.window_visits(s, np.array([o]), np.array([Lbar]), W)[0])
                    var_f += lam * v * v
            V = np.array([sum(K.visits_on_day(s, k) for k in range(L)) for L in pool], dtype=float)
            lamq = sum(B.adm[m][s].get(dt.date(2027, mo, 1), 0) for mo in (1, 2, 3))
            var_w += lamq * (V ** 2).mean()
        out["R3x_fixed_stays"][m] = B.plan_q[m] + z * math.sqrt(var_f)
        out["R3x_whole_stay_normal"][m] = B.plan_q[m] + z * math.sqrt(var_w)
    # a time-series reflex: AR(1) with a linear trend on daily visits since July 2024, the quarter's forecast-error
    # spread scaled to the plan
    out["R1_ar1_daily"] = {}
    for m in K.MEMBERS:
        y = ep[m].values.astype(float)
        t = np.arange(len(y), dtype=float)
        X = np.column_stack([np.ones(len(y) - 1), t[1:], y[:-1]])
        co, *_ = np.linalg.lstsq(X, y[1:], rcond=None)
        phi = float(co[2])
        sig2 = float(np.var(y[1:] - X @ co, ddof=3))
        h0 = (K.Q0 - K.D1).days
        h1 = (K.Q1 - K.D1).days
        w = np.array([sum(phi ** (h - k) for h in range(max(k, h0), h1 + 1)) for k in range(1, h1 + 1)])
        sd = math.sqrt(sig2 * float((w ** 2).sum()))
        mean_hist = float(y[-365:].mean()) * (h1 - h0 + 1)
        out["R1_ar1_daily"][m] = B.plan_q[m] + z * sd * B.plan_q[m] / mean_hist
    return {k: rung(B, v) for k, v in out.items()}


def dispersion(B):
    """Per Member, the variance-to-mean ratio of daily admissions around each month's own mean (episodes admitted
    1 April 2024 to 30 September 2026): near 1 is a Poisson arrival stream."""
    out = {}
    dd = K.days(K.D0, K.D1)
    for m in K.MEMBERS:
        cnt = {d: 0 for d in dd}
        for e in B.eps:
            if e["member"] == m and e["adm"] in cnt:
                cnt[e["adm"]] += 1
        by_month = {}
        for d in dd:
            by_month.setdefault((d.year, d.month), []).append(cnt[d])
        num = sum(float(np.sum((np.array(v) - np.mean(v)) ** 2)) for v in by_month.values())
        den = sum(len(v) - 1 for v in by_month.values())
        mu = float(np.mean([cnt[d] for d in dd]))
        out[m] = num / den / mu
    return out


def family(B, G=None):
    """The golden family: every admissible construction of the decisive step, which must book the same lots and round
    to the same requirement."""
    G = G or golden(B)
    fam = {}
    pools = {"pool_alliance_2024-04_to_2026-06": stays(B),
             "pool_post_revision_2025-07_to_2026-06": stays(B, first=K.REVISION),
             "pool_2025-01_to_2026-06": stays(B, first=dt.date(2025, 1, 1)),
             "pool_2024-10_to_2026-06": stays(B, first=dt.date(2024, 10, 1)),
             "pool_with_open_at_stay_to_date": stays(B, include_open=True, include_late=True),
             "pool_with_late_completed": stays(B, include_late=True),
             "pool_km_completed": stays(B, km=True)}
    for k, pool in pools.items():
        fam[k] = {m: requirement(B, m, pool)[0] for m in K.MEMBERS}
    fam["not_recentred"] = {m: requirement(B, m, recentre=False)[0] for m in K.MEMBERS}
    fam["start_2026-09-01"] = {m: requirement(B, m, start=dt.date(2026, 9, 1))[0] for m in K.MEMBERS}
    # per member stays where the member has them (the pathway members take the alliance's bone and joint stays)
    own = {}
    for m in K.MEMBERS:
        pm = stays(B, member=m)
        for s in K.STREAMS:
            if len(pm[s]) == 0:
                pm[s] = G["pool"][s]
        own[m] = requirement(B, m, pm)[0]
    fam["stays_per_member"] = own
    # interrupted courses merged, admissions reduced by the share that resumes a course
    share = continuation_share(B)
    pm = stays(B, merge=True)
    fam["courses_merged"] = {}
    for m in K.MEMBERS:
        adm = {s: {mo: v * (1 - share[s]) for mo, v in B.adm[m][s].items()} for s in K.STREAMS}
        fam["courses_merged"][m] = requirement(B, m, pm, adm=adm)[0]
    # the compound normal approximation
    fam["normal_approximation"] = {m: B.plan_q[m] + 1.2815516 * G["det"][m]["sd"] for m in K.MEMBERS}
    # a Monte Carlo simulation of 20,000 quarters
    fam["simulation_20000"] = {m: simulate_p90(B, m, G["pool"], 20000, seed=11)[0] for m in K.MEMBERS}
    return {k: rung(B, v) for k, v in fam.items()}


def simulate_p90(B, m, pool, nsim, seed=1, start=K.PATH):
    """A plain simulation: admissions each day at the plan's daily rate (Poisson), each with a stay drawn from the
    stream's stays, visits counted inside the quarter; the 90th percentile recentred on the plan."""
    r = np.random.default_rng(seed)
    dd = K.days(start, K.Q1)
    offs = np.array([(d - K.Q0).days for d in dd])
    W = (K.Q1 - K.Q0).days
    tot = np.zeros(nsim)
    for s in K.STREAMS:
        if len(pool[s]) == 0:
            continue
        rates = MD.daily_rates(B.adm[m][s], start, K.Q1)
        if rates.sum() == 0:
            continue
        n = r.poisson(np.tile(rates, (nsim, 1)))
        tn = n.sum(axis=1)
        day = np.concatenate([np.repeat(np.arange(len(dd)), n[j]) for j in range(nsim)])
        sid = np.repeat(np.arange(nsim), tn)
        L = r.choice(pool[s], len(day))
        v = MD.window_visits(s, offs[day], L, W)
        tot += np.bincount(sid, weights=v, minlength=nsim)
    p90 = float(np.percentile(tot, 90))
    return B.plan_q[m] + (p90 - tot.mean()), p90, tot.mean()


# ---------------------------------------------------------------------------------------------------- the asks

def load_asks(target):
    A = Bundle()
    p = lambda f: os.path.join(target, f)
    A.accounts = {r["client_code"]: r["alliance_member_code"] for r in _csv(p(K.ACCOUNTS))}
    A.holidays = {_date(r["date"]) for r in _csv(p(K.CALENDAR)) if r["public_holiday_vic"]}
    A.portal = _csv(p(K.PORTAL))
    wb = openpyxl.load_workbook(p(K.REGISTER), data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r[0] == "Log no.")
    head = rows[hi]
    A.register = [dict(zip(head, r)) for r in rows[hi + 1:] if r[0]]
    # the intake desk's file for a patient referred by telephone carries the log number until it is linked
    A.intake = {}
    for r in A.register:
        if r["Call type"] == "Referral" and r["Organisation"] in MEMBER_BY_ORG and r["Patient"]:
            m = MEMBER_BY_ORG[r["Organisation"]]
            A.intake[r["Log no."]] = (m, digits(r["Patient"]))
    A.visits = []
    for r in K.read_jsonl(p(K.BB_VISITS)):
        st = K.melbourne(dt.datetime.fromisoformat(r["start_time"]))
        en = K.melbourne(dt.datetime.fromisoformat(r["end_time"])) if r["end_time"] else None
        rec = K.melbourne(dt.datetime.fromisoformat(r["recorded_at"]))
        status = {"COMPLETE": "done", "Completed": "done", "NOTHOME": "attempt", "Attempted": "attempt",
                  "CANCELLED": "cancel", "Cancelled": "cancel"}[r["status"]]
        member = A.accounts[r["client_code"]]
        tail = r["patient_ref"].split("/", 1)[1]
        if not member and tail in A.intake:
            member = A.intake[tail][0]        # the patient's Member, from the intake log
        A.visits.append(dict(no=r["visit_no"], client=r["client_code"], member=member, pref=r["patient_ref"],
                             nurse=r["nurse_id"], type=r["visit_type"], start=st, end=en, status=status, rec=rec, raw=r))
    return A


KEEP = {"first_posted": (lambda v: (v["rec"], v["no"]), False),
        "last_posted": (lambda v: (v["rec"], v["no"]), True),
        "lowest_visit_no": (lambda v: v["no"], False),
        "highest_visit_no": (lambda v: v["no"], True),
        "earliest_start": (lambda v: (v["start"], v["rec"], v["no"]), False),
        "latest_start": (lambda v: (v["start"], v["rec"], v["no"]), True)}


def dedupe(visits, keep="first_posted"):
    """One record per visit: a record whose times overlap a kept record of the same nurse and patient is the same
    visit keyed again. `keep` names which record of such a pair survives (the golden keeps the first posted); the
    records come back in posting order whichever survives."""
    order, rev = KEEP[keep]
    kept = []
    by = {}
    for v in sorted(visits, key=order, reverse=rev):
        if v["status"] == "cancel":
            kept.append(v)
            continue
        key = (v["nurse"], v["pref"])
        dup = any(v["start"] < w["end"] and w["start"] < v["end"] for w in by.get(key, []))
        if not dup:
            by.setdefault(key, []).append(v)
            kept.append(v)
    return sorted(kept, key=lambda v: (v["rec"], v["no"]))


def fee(v, holidays):
    if v["status"] == "cancel":
        return 0.0
    if v["status"] == "attempt":
        return K.ATTEMPT_FEE
    t = v["start"]
    if t.date() in holidays:
        load = K.LOAD_HOLIDAY
    elif t.weekday() == 6:
        load = K.LOAD_SUN
    elif t.weekday() == 5:
        load = K.LOAD_SAT
    elif t.hour < 7 or t.hour >= 18:
        load = K.LOAD_AFTER_HOURS
    else:
        load = 0.0
    return K.FEE[v["type"]] * (1 + load)


def ask_a(A, visits=None, holidays=None):
    """What the alliance owes Bellbird for each member and month, July to September 2026, unrounded."""
    vs = dedupe(A.visits) if visits is None else visits
    hol = A.holidays if holidays is None else holidays
    out = {(m, mo): 0.0 for m in K.MEMBERS for mo in (7, 8, 9)}
    for v in vs:
        if not (K.Q3S <= v["start"].date() <= K.Q3E):
            continue
        out[(v["member"], v["start"].month)] += fee(v, hol)
    return out


WEEKS = [dt.date(2026, 6, 29) + dt.timedelta(days=7 * i) for i in range(14)]


def weekly(A, visits=None):
    """Bellbird's payable visits (completed or attempted) by member and week (weeks from Monday)."""
    vs = dedupe(A.visits) if visits is None else visits
    out = {(m, w): 0 for m in K.MEMBERS for w in WEEKS}
    for v in vs:
        if v["status"] == "cancel" or not (K.Q3S <= v["start"].date() <= K.Q3E):
            continue
        d = v["start"].date()
        w = d - dt.timedelta(days=d.weekday())
        out[(v["member"], w)] += 1
    return out


MEMBER_BY_ORG = {"Northgate Health": "NGH", "Riverbend Health": "RBH", "Westmere Health": "WMH",
                 "Westmere Health (Altona Park)": "WMH", "Kingsdale Health": "KDH", "Hillcrest Health": "HCH"}


def digits(s):
    return "".join(ch for ch in str(s) if ch.isdigit()).lstrip("0")


def referrals(A, register=True):
    """Every referral a member made to Bellbird in July to September 2026: the portal's, and those telephoned to
    Bellbird's intake line (the register), each with its time, outcome, response time and patient key."""
    out = []
    for r in A.portal:
        sent = dt.datetime.strptime(r["sent_at"], "%Y-%m-%d %H:%M")
        if not (K.Q3S <= sent.date() <= K.Q3E):
            continue
        resp = K.melbourne(dt.datetime.strptime(r["bellbird_response_at"], "%Y-%m-%dT%H:%M:%SZ"))
        out.append(dict(member=r["member_code"], sent=sent, accepted=r["status"] == "Accepted", resp=resp,
                        key=(r["member_code"], digits(r["patient_no"])), source="portal"))
    if register:
        for r in A.register:
            if r["Call type"] != "Referral" or r["Organisation"] not in MEMBER_BY_ORG:
                continue
            d = r["Date"].date() if isinstance(r["Date"], dt.datetime) else r["Date"]
            if not (K.Q3S <= d <= K.Q3E):
                continue
            sent = dt.datetime.combine(d, dt.datetime.strptime(r["Time"], "%H:%M").time())
            resp = dt.datetime.combine(d, dt.datetime.strptime(r["Outcome time"], "%H:%M").time()) \
                if r["Outcome time"] else None
            m = MEMBER_BY_ORG[r["Organisation"]]
            out.append(dict(member=m, sent=sent, accepted=r["Outcome"] == "Accepted", resp=resp,
                            key=(m, digits(r["Patient"])), source="register"))
    return out


def patient_key(A, v, link=True):
    """(member, patient digits) of a Bellbird record; an intake-file reference resolves through the register."""
    tail = v["pref"].split("/", 1)[1]
    if tail.startswith("IL26-"):
        return A.intake[tail] if link else (v["member"], tail)
    return (v["member"], digits(tail))


def ask_b(A, refs=None, visits=None, link=True):
    """Per member: referrals made, per cent accepted within two hours, per cent of accepted referrals first visited
    within 24 hours of acceptance (unrounded), with the counts."""
    refs = referrals(A) if refs is None else refs
    vs = dedupe(A.visits) if visits is None else visits
    att = {}
    for v in vs:
        if v["status"] == "cancel":
            continue
        att.setdefault(patient_key(A, v, link=link), []).append(v["start"])
    for k in att:
        att[k].sort()
    out = {}
    for m in K.MEMBERS:
        rr = [x for x in refs if x["member"] == m]
        n = len(rr)
        acc2 = sum(1 for x in rr if x["accepted"] and (x["resp"] - x["sent"]).total_seconds() <= 2 * 3600)
        acc = [x for x in rr if x["accepted"]]
        v24 = 0
        lags = []
        for x in acc:
            ts = [t for t in att.get(x["key"], []) if t >= x["resp"]]
            if ts:
                lag = (ts[0] - x["resp"]).total_seconds() / 3600
                lags.append(lag)
                if lag <= K.FIRST_VISIT_HOURS:
                    v24 += 1
        out[m] = dict(n=n, acc2=acc2, acc=len(acc), v24=v24, pct_acc2=100.0 * acc2 / n,
                      pct_v24=100.0 * v24 / len(acc), lags=lags,
                      resp_lags=[(x["resp"] - x["sent"]).total_seconds() / 60 for x in rr if x["accepted"]])
    return out
