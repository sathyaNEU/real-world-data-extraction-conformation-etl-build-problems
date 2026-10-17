"""task193: the governing standard as a PDF and the statutory fee schedule as a workbook."""
import os

from reportlab import rl_config
rl_config.invariant = 1
from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                Table, TableStyle)

import world as W

SS = getSampleStyleSheet()
BODY = ParagraphStyle("b", parent=SS["BodyText"], fontName="Times-Roman", fontSize=9.5, leading=13,
                      alignment=TA_JUSTIFY, spaceAfter=5)
H1 = ParagraphStyle("h1", parent=SS["Heading1"], fontName="Times-Bold", fontSize=13, spaceAfter=8)
H2 = ParagraphStyle("h2", parent=SS["Heading2"], fontName="Times-Bold", fontSize=10.5, spaceAfter=4,
                    spaceBefore=8)
SMALL = ParagraphStyle("s", parent=BODY, fontSize=8, leading=10.5)
CELL = ParagraphStyle("c", fontName="Times-Roman", fontSize=8, leading=10)
CELLB = ParagraphStyle("cb", fontName="Times-Bold", fontSize=8, leading=10)
TBL = TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.Color(0.45, 0.45, 0.45)),
                  ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.90, 0.90, 0.88)),
                  ("VALIGN", (0, 0), (-1, -1), "TOP"),
                  ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                  ("TOPPADDING", (0, 0), (-1, -1), 1.6), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6)])


def _table(head, rows, widths):
    data = [[Paragraph(h, CELLB) for h in head]]
    for r in rows:
        data.append([Paragraph(str(c), CELL) for c in r])
    t = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
    t.setStyle(TBL)
    return t


def standard(tgt):
    path = os.path.join(tgt, "screen_standard_7_2024.pdf")
    doc = SimpleDocTemplate(path, pagesize=A4, title="Standard 7-2024", author="", subject="",
                            creator="", leftMargin=22 * mm, rightMargin=22 * mm,
                            topMargin=20 * mm, bottomMargin=18 * mm)
    f = []
    f.append(Paragraph("Standard 7-2024", H1))
    f.append(Paragraph("The quarterly prescription-sequence screen of medicines under additional "
                       "monitoring", H2))
    f.append(Paragraph("Issued by the pharmacovigilance division. Version 7-2024.3, in force from "
                       "1 July 2025. Supersedes 7-2024.2.", SMALL))
    f.append(Paragraph("1. Purpose and scope", H2))
    f.append(Paragraph("1.1 This standard governs the quarterly screen the division runs on the "
                       "Register of Medicinal Product Statistics for medicines under additional "
                       "monitoring. The screen looks for pairs of a monitored product and a marker "
                       "treatment where the marker treatment is started after the product more "
                       "often than the register's own dispensing pattern would give, and refers "
                       "such pairs for formal signal assessment.", BODY))
    f.append(Paragraph("1.2 The run's extract, its screen window and the products and marker "
                       "treatments screened are fixed for each run and recorded with it.", BODY))

    f.append(Paragraph("2. The run", H2))
    f.append(Paragraph("2.1 The screen window. Each quarterly run covers the twelve months ending "
                       "on the last day of the quarter before the run, both days included. A "
                       "dispensing falls in the screen window, and in a quarter of it, according "
                       "to its date of dispensing and not according to the date it was submitted "
                       "to the register.", BODY))
    f.append(Paragraph("2.2 The clearance period. The twenty-four months before the screen window, "
                       "that is 1 October 2023 to 30 September 2025, are the clearance period. The "
                       "register extract for a run covers the clearance period and the screen "
                       "window together.", BODY))
    f.append(Paragraph("2.3 A first receipt. A patient's first receipt of a treatment is the "
                       "earliest dispensing to that patient, anywhere in the extract, of any "
                       "package whose anatomical therapeutic chemical code falls in that "
                       "treatment's group. A patient whose first receipt of a treatment falls in "
                       "the clearance period has no first receipt of that treatment in the screen "
                       "window.", BODY))
    f.append(Paragraph("2.4 One patient, one medicine. A person key replaced inside the extract "
                       "stands for the same patient before and after the replacement. A package "
                       "code belongs to the product the package register gives for it on the date "
                       "of the dispensing.", BODY))
    f.append(Paragraph("2.5 A pair's patients. The run screens each pair of a product in Annex A "
                       "and a marker treatment in Annex B. A pair's patients are the patients "
                       "whose first receipt of the product and whose first receipt of the marker "
                       "treatment both fall inside the screen window. A patient whose two first "
                       "receipts fall on the same day is not one of the pair's patients. Each "
                       "patient counts once in each pair.", BODY))

    f.append(Paragraph("3. The referral test", H2))
    f.append(Paragraph("3.1 The order. For each of a pair's patients the run records whether the "
                       "first receipt of the marker treatment fell after the first receipt of the "
                       "product, or before it.", BODY))
    f.append(Paragraph("3.2 Referral. A pair is referred for formal signal assessment where", BODY))
    f.append(Paragraph("(a) the pair has at least 25 patients;", BODY))
    f.append(Paragraph("(b) the number of the pair's patients in whom the first receipt of the "
                       "marker treatment fell after the first receipt of the product is at least "
                       "1.25 times the number that would be expected if receiving the product had "
                       "no bearing on when the marker treatment is first received; and", BODY))
    f.append(Paragraph("(c) the pair is not recorded as under assessment in the signal management "
                       "tracker at the date the extract was pulled.", BODY))
    f.append(Paragraph("3.3 Figures. Patient counts are whole. The expected number is carried to "
                       "one decimal and the multiple in 3.2(b) to two decimals.", BODY))
    f.append(Paragraph("3.4 A pair that meets 3.2(a) but not 3.2(b) is carried forward to the next "
                       "run unchanged and is not reported as a referral.", BODY))
    f.append(Paragraph("3.5 The construction of the expected number in 3.2(b) is the division's "
                       "own epidemiological work and is recorded with each run.", BODY))

    f.append(Paragraph("4. After the run", H2))
    f.append(Paragraph("4.1 The referral quarter. Pairs a run refers are assessed in the quarter "
                       "that begins after the quarter the run is made in. The division publishes "
                       "the run's figures for every screened pair, referred or not.", BODY))
    f.append(Paragraph("4.2 The standing utilisation table. The division's quarterly table counts "
                       "every patient for whom at least one product in Annex A has its first "
                       "receipt inside the screen window, and gives the share of those patients "
                       "for whom at least one marker treatment in Annex B also has its first "
                       "receipt inside the screen window. Patients are counted by the region "
                       "recorded for them at the start of the screen window and by their age at "
                       "that date, taken as 2025 less the year of birth the register holds, in the "
                       "bands 18 to 39, 40 to 59, 60 to 74, and 75 and over. A patient for whom no "
                       "region is recorded is counted under \"not stated\".", BODY))

    f.append(PageBreak())
    rows = []
    for code, brand, atc, prof, pi, mf, mt, holder in W.PRODUCTS:
        rows.append([code, brand, atc, ", ".join(W.ANNEX[code])])
    f.append(KeepTogether([
        Paragraph("Annex A. Products under additional monitoring during the screen window", H2),
        _table(["Product code", "Brand name", "ATC code", "Marker groups screened"],
               rows, [22, 30, 24, 44])]))
    f.append(Spacer(1, 6))
    rows = [[m[0], m[1], m[2]] for m in W.MARKERS]
    f.append(KeepTogether([
        Paragraph("Annex B. Marker treatments", H2),
        _table(["Marker code", "Marker treatment", "ATC group"], rows, [24, 54, 30])]))
    f.append(Spacer(1, 4))
    f.append(Paragraph("A marker treatment is the treatment a clinician gives for the event the "
                       "division is watching for, and the division screens for the treatment "
                       "because the register records it while it does not record the event. A "
                       "package belongs to a marker treatment where its anatomical therapeutic "
                       "chemical code is the group code in this annex.", SMALL))
    doc.build(f)


def fee_schedule(tgt, pulled):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Schedule"
    head_font = Font(name="Calibri", size=11, bold=True)
    fill = PatternFill("solid", fgColor="DCE6F1")
    ws["A1"] = "Statutory data fee for the prescription-sequence screen"
    ws["A1"].font = Font(name="Calibri", size=12, bold=True)
    ws["A2"] = "Executive order 1184 of 29 August 2024, schedule 3. Rates in DKK per 1.000 DDD."
    ws["A3"] = ("The fee is charged on defined daily doses dispensed inside the screen window, "
                "whatever month the line was submitted in; a line submitted after the quarter "
                "closed is charged in the quarter of its dispensing.")
    ws["A4"] = "Doses are stated in thousands of DDD on this sheet."
    ws["A5"] = ("A payer's band is settled once, on its total doses for the whole screen window. "
                "The band's rate is charged on all of those doses; the bands are not applied in "
                "tiers. Doses dispensed before 1 April 2026 are charged at this sheet's rate for "
                "the band and doses dispensed on or after that date at the Amendment sheet's rate "
                "for the same band.")
    rows = [("Band", "From (thousand DDD)", "To (thousand DDD)", "Rate (DKK per 1.000 DDD)")]
    for i, (lo, hi, rt) in enumerate(W.FEE_BANDS):
        rows.append(("%d" % (i + 1), "%.0f" % lo, "-" if hi > 10 ** 8 else "%.0f" % hi, "%d" % rt))
    for r, row in enumerate(rows, start=7):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            if r == 7:
                cell.font = head_font
                cell.fill = fill
    ws2 = wb.create_sheet("Amendment 2026-04-01")
    ws2["A1"] = "Amendment to schedule 3, in force from 1 April 2026"
    ws2["A1"].font = Font(name="Calibri", size=12, bold=True)
    ws2["A2"] = ("Executive order 288 of 11 March 2026. The rates below apply to defined daily "
                 "doses dispensed on or after 1 April 2026. Doses dispensed before that date stay "
                 "on the rates in the Schedule sheet.")
    ws2["A3"] = "Doses are stated in thousands of DDD on this sheet."
    rows = [("Band", "From (thousand DDD)", "To (thousand DDD)", "Rate (DKK per 1.000 DDD)")]
    for i, (lo, hi, rt) in enumerate(W.FEE_BANDS_AMENDED):
        rows.append(("%d" % (i + 1), "%.0f" % lo, "-" if hi > 10 ** 8 else "%.0f" % hi, "%d" % rt))
    for r, row in enumerate(rows, start=5):
        for c, v in enumerate(row, start=1):
            cell = ws2.cell(row=r, column=c, value=v)
            if r == 5:
                cell.font = head_font
                cell.fill = fill
    ws3 = wb.create_sheet("Notes")
    ws3["A1"] = "Notes"
    ws3["A1"].font = Font(name="Calibri", size=12, bold=True)
    ws3["A2"] = ("The fee payer is the marketing authorisation holder of the product on the date of "
                 "the dispensing, as the product register gives it for that date. Where a holder "
                 "has merged into another company, the surviving company is the fee payer for "
                 "dispensings on and after the merger date, and the merged company remains the "
                 "payer for its dispensings before that date.")
    ws3["A3"] = ("A dispensing re-transmitted after a failed transfer is charged once. Every "
                 "dispensing inside the screen window is charged, including a dispensing of a "
                 "product whose marketing authorisation has since lapsed.")
    ws3["A4"] = "Prepared by the finance office, " + W.iso(pulled) + "."
    for w in (ws, ws2, ws3):
        for col in range(1, 6):
            w.column_dimensions[get_column_letter(col)].width = [10, 24, 24, 28, 14][col - 1]
        w.sheet_view.showGridLines = False
        for row in w.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and len(cell.value) > 60:
                    cell.alignment = Alignment(wrap_text=False, vertical="top")
    wb.save(os.path.join(tgt, "data_fee_schedule.xlsx"))
