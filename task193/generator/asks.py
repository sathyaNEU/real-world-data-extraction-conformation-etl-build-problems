"""The pipeline's own files: the submissions the agency could not code, the channel totals, the
holders' submissions and their data agreements. None of them carries a coded reaction."""
import csv
import datetime as dt
import os
import random

import world as W
from pack import wr

# the holders' data agreements: rate per submission in DKK, and the one amended inside the quarter
RATES = {"Aurelind Pharma A/S": 148, "Brynholm Laegemidler ApS": 132, "Castellan Nordic A/S": 165,
         "Dalberg Therapeutics A/S": 121, "Elvershoj Biotek A/S": 154,
         "Fjeldgaard Pharma A/S": 139, "Granvig Medical A/S": 118,
         "Hvidsten Laboratorier A/S": 127, "Nordsund Vaccines A/S": 96,
         "Kvistgaard Onkologi A/S": 154}
AMENDED = ("Castellan Nordic A/S", 165, 184, dt.date(2026, 8, 1))
# one holder files under two codes after a merger inside the quarter
MERGED = ("Hvidsten Laboratorier A/S", "MAH-0471", "MAH-0518", dt.date(2026, 8, 24))
HOLDER_CODE = {h: "MAH-%04d" % (380 + 17 * i) for i, h in enumerate(sorted(RATES))}
HOLDER_CODE[MERGED[0]] = MERGED[1]

HOLDER_OF = {}
for _c, _n, _h, _a, _k, _e in W.PRODUCTS:
    HOLDER_OF[_c] = _h
HOLDER_OF[W.DOMINANT] = "Nordsund Vaccines A/S"
HOLDER_OF[W.RENAMED[0]] = W.RENAMED[4]

FAIL_REASON = ["product not identifiable", "no reaction stated", "duplicate of an open case",
               "reporter contact missing", "term not in the dictionary"]


def exceptions(cases):
    """submissions the agency received and could not code, so they never reach the extract"""
    rng = random.Random(W.SEED * 43 + 17)
    months = [(dt.date(2026, 7, 1), dt.date(2026, 7, 31)),
              (dt.date(2026, 8, 1), dt.date(2026, 8, 31)),
              (dt.date(2026, 9, 1), dt.date(2026, 9, 30)),
              (dt.date(2026, 10, 1), W.CODED_BY)]
    per = {"clinician": (14, 23), "licence holder": (11, 19), "patient": (9, 16),
           "pharmacy": (4, 8)}
    out = []
    sid = 1900000
    for lo, hi in months:
        span = (hi - lo).days
        for ch in W.CHANNELS:
            a, b = per[ch]
            n = rng.randint(a, b) if lo.month != 10 else rng.randint(2, 5)
            for _ in range(n):
                sid += 1
                d = lo + dt.timedelta(days=rng.randint(0, span))
                out.append({"sub": "SUB-%07d" % sid, "channel": ch, "receipt": d,
                            "stamp": dt.datetime.combine(d, dt.time(rng.randrange(7, 19),
                                                                    rng.randrange(0, 60))),
                            "holder": rng.choice(sorted(RATES)) if ch == "licence holder" else "",
                            "reason": rng.choice(FAIL_REASON)})
    out.sort(key=lambda e: (e["receipt"], e["sub"]))
    return out


def write_exceptions(tgt, exc):
    rows = [[e["sub"], e["channel"], e["receipt"].isoformat(),
             e["stamp"].strftime("%Y-%m-%dT%H:%M:00+02:00"), e["holder"], e["reason"]]
            for e in exc]
    return wr(os.path.join(tgt, "coding_exceptions.csv"),
              ["submission_id", "channel", "received_on", "submitted_at_sender_stamp",
               "licence_holder", "reason_not_coded"], rows)


def write_channel_totals(tgt, subs, exc):
    """the agency's own daily count of what each channel sent, coded or not"""
    tally = {}
    for s in subs:
        tally[(s["receipt"], s["channel"])] = tally.get((s["receipt"], s["channel"]), 0) + 1
    for e in exc:
        tally[(e["receipt"], e["channel"])] = tally.get((e["receipt"], e["channel"]), 0) + 1
    rows = [[d.isoformat(), ch, n] for (d, ch), n in sorted(tally.items())]
    return wr(os.path.join(tgt, "channel_totals.csv"),
              ["received_on", "channel", "submissions_received"], rows)


def write_holder_submissions(tgt, subs, exc):
    """every submission a licence holder filed, under the holder code in force on the day"""
    rows = []
    for s in subs:
        if s["channel"] != "licence holder":
            continue
        h = HOLDER_OF[s["code"]]
        code = HOLDER_CODE[h]
        if h == MERGED[0] and s["receipt"] >= MERGED[3]:
            code = MERGED[2]
        rows.append([s["sub"], code, h, s["receipt"].isoformat(), "coded"])
    for e in exc:
        if e["channel"] != "licence holder" or not e["holder"]:
            continue
        code = HOLDER_CODE[e["holder"]]
        if e["holder"] == MERGED[0] and e["receipt"] >= MERGED[3]:
            code = MERGED[2]
        rows.append([e["sub"], code, e["holder"], e["receipt"].isoformat(), "not coded"])
    rows.sort(key=lambda r: (r[3], r[0]))
    return wr(os.path.join(tgt, "holder_submissions.csv"),
              ["submission_id", "holder_code", "licence_holder", "received_on", "coding_status"],
              rows)


def write_agreements(tgt):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    ws["A1"] = "Data agreement schedule, reporting fees payable to licence holders"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Rates are per submission received and coded or not coded, as the agreement provides."
    ws["A2"].font = Font(size=9, italic=True)
    ws.append([])
    ws.append(["Holder code", "Licence holder", "Rate per submission (DKK)", "In force from",
               "In force to"])
    for c in ws[4]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for h in sorted(RATES):
        code = HOLDER_CODE[h]
        if h == AMENDED[0]:
            ws.append([code, h, AMENDED[1], "2025-01-01", (AMENDED[3] - dt.timedelta(days=1)).isoformat()])
        else:
            ws.append([code, h, RATES[h], "2025-01-01", None])
    ws2 = wb.create_sheet("Amendments")
    ws2["A1"] = "Amendments to the schedule"
    ws2["A1"].font = Font(bold=True, size=12)
    ws2.append([])
    ws2.append(["Holder code", "Licence holder", "Amendment", "Rate per submission (DKK)",
                "In force from", "Signed"])
    for c in ws2[3]:
        c.font = Font(bold=True)
    ws2.append([HOLDER_CODE[AMENDED[0]], AMENDED[0], "AGR-2026-041", AMENDED[2],
                AMENDED[3].isoformat(), "2026-07-21"])
    ws3 = wb.create_sheet("Holder codes")
    ws3.append(["Holder code", "Licence holder", "In force from", "Note"])
    for c in ws3[1]:
        c.font = Font(bold=True)
    for h in sorted(RATES):
        ws3.append([HOLDER_CODE[h], h, "2025-01-01", None])
    ws3.append([MERGED[2], MERGED[0], MERGED[3].isoformat(),
                "code reissued on the merger of 24 August 2026"])
    for w in (ws, ws2, ws3):
        for i, wid in enumerate((14, 34, 26, 16, 16, 14), start=1):
            w.column_dimensions[chr(64 + i)].width = wid
    p = os.path.join(tgt, "data_agreement_schedule.xlsx")
    wb.save(p)
    return p


def write_distractors(tgt):
    rng = random.Random(W.SEED * 47 + 19)
    rows = []
    d = W.Q_FROM
    JOURNALS = ["Nordisk Farmakologi", "Acta Clinica Scandinavica", "European Drug Safety Review",
                "Journal of Pharmacoepidemiology", "Klinisk Praksis"]
    while d <= W.Q_TO:
        for _ in range(rng.choices([0, 1, 2, 3], weights=[22, 38, 28, 12])[0]):
            rows.append([d.isoformat(), rng.choice(JOURNALS),
                         "REF-%05d" % rng.randrange(10000, 99999),
                         rng.choice(["screened out", "screened out", "screened out",
                                     "full text requested", "case series noted"]),
                         rng.choice(W.REGIONS + ["outside Denmark", "outside Denmark"])])
        d += dt.timedelta(days=1)
    n1 = wr(os.path.join(tgt, "literature_screening_log.csv"),
            ["screened_on", "journal", "reference", "outcome", "origin"], rows)
    rows = []
    d = W.Q_FROM
    while d <= W.Q_TO:
        for _ in range(rng.choices([2, 3, 4, 5, 6, 7], weights=[8, 16, 24, 22, 18, 12])[0]):
            rows.append([d.isoformat(),
                         rng.choice(["telephone", "telephone", "email", "web form"]),
                         rng.choice(["member of the public", "pharmacist", "physician", "nurse",
                                     "journalist"]),
                         rng.choice(["interaction", "dose", "storage", "reporting a reaction",
                                     "pregnancy", "availability"]),
                         rng.choices([2, 4, 6, 9, 13, 21], weights=[22, 26, 20, 16, 10, 6])[0]])
        d += dt.timedelta(days=1)
    n2 = wr(os.path.join(tgt, "enquiry_log.csv"),
            ["enquiry_on", "route", "caller", "subject", "minutes_on_call"], rows)
    return n1, n2
