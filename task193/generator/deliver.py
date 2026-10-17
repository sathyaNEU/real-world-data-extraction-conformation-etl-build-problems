"""task193: the two golden deliverables, the screen workbook and the order-share chart."""
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter


def workbook(path, g):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    thin = Side(style="thin", color="BFBFBF")
    head = Font(name="Calibri", size=11, bold=True)
    fill = PatternFill("solid", fgColor="E8EDF3")
    note = Font(name="Calibri", size=9, italic=True)

    ws = wb.active
    ws.title = "Screen 2026 Q4 run"
    ws["A1"] = "Prescription-sequence screen, run of October 2026"
    ws["A1"].font = Font(name="Calibri", size=13, bold=True)
    ws["A2"] = ("Screen window 1 October 2025 to 30 September 2026. Extract pulled 20 October 2026. "
                "Referrals are assessed in the quarter beginning 1 January 2027.")
    ws["A2"].font = note
    ws["A3"] = ("Expected after-count on the window's own first-receipt dates, Standard 7-2024 "
                "section 3.2(b); referral multiple 1.25; minimum 25 patients.")
    ws["A3"].font = note
    cols = ["Product", "Brand", "Marker treatment", "Patients", "Marker after", "Marker before",
            "Share after", "Expected after", "Multiple", "Referred"]
    for c, v in enumerate(cols, start=1):
        cell = ws.cell(row=5, column=c, value=v)
        cell.font = head
        cell.fill = fill
        cell.border = Border(bottom=thin)
    r = 6
    for row in g["screen_rows"]:
        vals = [row["product"], row["brand"], row["marker_treatment"], row["patients"],
                row["after"], row["before"],
                None if row["share_pct"] is None else row["share_pct"] / 100.0,
                row["expected"], row["multiple"], row["referred"]]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            if c == 7 and v is not None:
                cell.number_format = "0.0%"
            if c == 8 and v is not None:
                cell.number_format = "0.0"
            if c == 9 and v is not None:
                cell.number_format = "0.00"
        r += 1
    ws.cell(row=r + 1, column=1, value="Referred for assessment in 2027 Q1").font = head
    ws.cell(row=r + 1, column=4, value=g["referred_count"]).font = head
    ws.cell(row=r + 2, column=1,
            value=("Pairs below the 25-patient minimum are shown with no figures and are carried "
                   "forward to the next run under section 3.4.")).font = note
    ws.cell(row=r + 3, column=1,
            value=("Nearest pair on the wrong side of the line: %s, multiple %.2f."
                   % (g["nearest_miss"]["label"], g["nearest_miss"]["multiple"]))).font = note

    ws2 = wb.create_sheet("Chain submissions")
    ws2["A1"] = "Dispensing lines by chain and quarter of dispensing"
    ws2["A1"].font = Font(name="Calibri", size=12, bold=True)
    ws2["A2"] = ("Quarters are quarters of dispensing. Accepted lines are the distinct dispensings "
                 "in the two register extracts, counting a re-sent line once; lines turned back at "
                 "validation are in submission_exceptions.csv and never reach the extracts.")
    ws2["A2"].font = note
    for c, v in enumerate(["Chain", "Quarter", "Lines accepted", "Lines rejected",
                           "Median days to submission"], start=1):
        cell = ws2.cell(row=4, column=c, value=v)
        cell.font = head
        cell.fill = fill
    for i, row in enumerate(g["a1"], start=5):
        for c, v in enumerate([row["chain"], row["quarter"], row["accepted"], row["rejected"],
                               row["median_days"]], start=1):
            cell = ws2.cell(row=i, column=c, value=v)
            if c == 5:
                cell.number_format = "0.0"

    ws3 = wb.create_sheet("Holder fees")
    ws3["A1"] = "Statutory data fee, screen window 2025-26"
    ws3["A1"].font = Font(name="Calibri", size=12, bold=True)
    ws3["A2"] = ("Doses at the rates in schedule 3, with the amendment of 1 April 2026 applied to "
                 "doses dispensed on or after that date, and each payer's band rate charged on all "
                 "of its doses. Sovraxil's doses move to Ellevang with the authorisation on "
                 "27 February 2026; Strandby's own doses before the merger of 1 April 2026 are "
                 "charged to Strandby and its later doses to Bregnholt. Re-sent lines are counted "
                 "once.")
    ws3["A2"].font = note
    for c, v in enumerate(["Holder", "Name", "DDD dispensed", "Fee due (DKK)"], start=1):
        cell = ws3.cell(row=4, column=c, value=v)
        cell.font = head
        cell.fill = fill
    for i, row in enumerate(g["a2"], start=5):
        for c, v in enumerate([row["holder"], row["name"], row["ddd"], row["fee_dkk"]], start=1):
            cell = ws3.cell(row=i, column=c, value=v)
            if c == 3:
                cell.number_format = "#,##0.0"
            if c == 4:
                cell.number_format = "#,##0"

    ws4 = wb.create_sheet("Utilisation table")
    ws4["A1"] = "Patients with a first receipt of a screened product in the window"
    ws4["A1"].font = Font(name="Calibri", size=12, bold=True)
    ws4["A2"] = ("Region and age as recorded at 1 October 2025, Standard 7-2024 section 4.2. "
                 "Share is of the same patients who also had a first marker-treatment receipt in "
                 "the window.")
    ws4["A2"].font = note
    for c, v in enumerate(["Region", "Age band", "Patients", "Share with a first marker receipt"],
                          start=1):
        cell = ws4.cell(row=4, column=c, value=v)
        cell.font = head
        cell.fill = fill
    for i, row in enumerate(g["a3"], start=5):
        for c, v in enumerate([row["region"], row["band"], row["patients"],
                               row["share_pct"] / 100.0], start=1):
            cell = ws4.cell(row=i, column=c, value=v)
            if c == 4:
                cell.number_format = "0.0%"

    widths = {"Screen 2026 Q4 run": [10, 13, 23, 10, 13, 14, 11, 15, 10, 10],
              "Chain submissions": [9, 11, 15, 15, 26],
              "Holder fees": [9, 29, 16, 16],
              "Utilisation table": [15, 13, 11, 32]}
    for name, ws_ in (("Screen 2026 Q4 run", ws), ("Chain submissions", ws2),
                      ("Holder fees", ws3), ("Utilisation table", ws4)):
        for c, w in enumerate(widths[name], start=1):
            ws_.column_dimensions[get_column_letter(c)].width = w
        ws_.sheet_view.showGridLines = False
        ws_.freeze_panes = "A6" if name == "Screen 2026 Q4 run" else "A5"
    wb.save(path)


def chart(path, g):
    rows = [r for r in g["screen_rows"] if r["share_pct"] is not None]
    rows = sorted(rows, key=lambda r: r["share_pct"])
    labels = ["%s + %s" % (r["brand"], r["marker_treatment"]) for r in rows]
    vals = [r["share_pct"] for r in rows]
    REF, NOT = "#1f4e79", "#aebfd2"
    LINE = "#a4161a"
    cols = [REF if r["referred"] == "yes" else NOT for r in rows]
    fig, ax = plt.subplots(figsize=(10.4, 7.4), dpi=150)
    y = list(range(len(rows)))
    ax.barh(y, vals, color=cols, height=0.66, edgecolor="white", linewidth=0.9)
    ax.set_yticks(y)
    ax.set_yticklabels(labels, fontsize=9)
    ax.set_xlabel("Share of the pair's patients who received the marker treatment after the product",
                  fontsize=9.5, labelpad=8)
    ax.set_xlim(0, 108)
    ax.set_xticks([0, 20, 40, 60, 80, 100])
    ax.xaxis.set_major_formatter(FuncFormatter(lambda v, p: "%d%%" % v))
    ax.tick_params(axis="x", labelsize=9)
    ax.tick_params(axis="y", length=0)
    for i, r in enumerate(rows):
        ax.text(r["share_pct"] + 1.1, i, "%.1f%%" % r["share_pct"], va="center", fontsize=8,
                color="#2b2b2b")
    from matplotlib.patches import Patch
    handles = [Patch(facecolor=REF, label="referred to assessment"),
               Patch(facecolor=NOT, label="not referred")]
    ax.legend(handles=handles, loc="lower center", bbox_to_anchor=(0.5, -0.135), ncol=2,
              fontsize=8.5, frameon=False, handlelength=1.4, columnspacing=1.6)
    ax.set_title("%d of the %d candidate pairs clear the referral line once the expected share is "
                 "taken\nfrom the window's own first-receipt dates, against %d on an equal-chance "
                 "expectation" % (g["referred_count"], g["candidate_count"], g["crude_count"]),
                 fontsize=11.5, loc="left", pad=34, color="#13344f")
    ax.text(0.0, 1.018,
            "Prescription-sequence screen, run of October 2026. Screen window 1 October 2025 to "
            "30 September 2026.", transform=ax.transAxes, fontsize=8.5, color="#5a5a5a")
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.spines["bottom"].set_color("#9a9a9a")
    ax.grid(axis="x", color="#e4e4e4", linewidth=0.7)
    ax.set_axisbelow(True)
    fig.text(0.012, 0.012,
             "Source: Laegemiddelstatistikregisteret extract LSR-2026-0441, pulled 20 October 2026. "
             "Standard 7-2024 sections 2.5, 3.1 and 3.2.", fontsize=7.6, color="#6a6a6a")
    fig.subplots_adjust(left=0.235, right=0.985, top=0.868, bottom=0.145)
    fig.savefig(path, metadata={"Software": None, "Creation Time": None})
    plt.close(fig)
