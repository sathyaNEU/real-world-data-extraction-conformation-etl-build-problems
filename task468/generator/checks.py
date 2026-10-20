"""The generator's assertions on the cut pack and the clean truth: the index population, the ladder (every route's set,
the golden and the readings that converge on it), the analyst test on the clean truth, the Layer 1 items left neutral on
the call, the margins, the semantic fork table, the bonus figures and their device, the separation of the ask from the
call, the input gates and the documents."""
import json
import os
import re
import subprocess

import numpy as np
import openpyxl
import pandas as pd

import bundle as B
import methods as M
import sim as S
import world as W

GOLD = sorted(["Calloway Memorial Hospital", "Mercy Valley Medical Center", "Northgate Community Hospital",
               "Brannock Medical Center", "Westover Regional Hospital"])
REFLEX = sorted(["Varden Surgical Hospital", "Calloway Memorial Hospital", "Northgate Community Hospital",
                 "Westover Regional Hospital"])
ROUTE_SETS = {   # every wrong route of the ladder and the set it files on the pack
    "naive": REFLEX, "union": REFLEX, "reflex": REFLEX,
    "fp_only": ["Calloway Memorial Hospital", "Northgate Community Hospital", "Westover Regional Hospital"],
    "recall_only": ["Varden Surgical Hospital"] + GOLD,
}
BY_SETS = {
    "pass": ["Varden Surgical Hospital", "Calloway Memorial Hospital", "Mercy Valley Medical Center",
             "Northgate Community Hospital", "Westover Regional Hospital"],
    "all": REFLEX,
}
DECISIVE_WORDS = ["false link", "false-link", "recall", "sensitivity", "specificity", "positive predictive",
                  "misclassif", "linkage error", "error rate", "two-part", "compound", "hyphen", "undercount",
                  "overcount", "missed death", "control group", "capture", "correct for", "corrected for"]


class Run:
    def __init__(self):
        self.n, self.bad = 0, []

    def check(self, name, cond, detail=""):
        self.n += 1
        if not cond:
            self.bad.append("%s %s" % (name, detail))
            print("FAIL", name, detail)


def pdf_text(p):
    return subprocess.run(["pdftotext", "-layout", p, "-"], capture_output=True, text=True).stdout


def run(target, golden_dir, x=None, meta_path=None, fig_path=None):
    R = Run()
    P = B.Pack(target)
    fig = json.load(open(fig_path, encoding="utf-8"))
    f = P.frame()
    g = P.determination(frame=f)

    # ---------------------------------------------------------------- the population and the golden
    x = x or __import__("pack").everything()
    truth_n = {h: sum(1 for s in x["stays"] if s["hosp"] == h) for h in W.CANDIDATES}
    R.check("P1 index admissions per hospital equal the clean truth", g["n"] == truth_n, str(g["n"]))
    R.check("SF4 every index admission has a HIMM 3.2 score", f["p"].notna().all())
    R.check("P3 the golden set", g["losers"] == GOLD, str(g["losers"]))
    R.check("P4 figures.json carries the golden set", fig["lose_full"] == GOLD and fig["count_lose"] == 5)
    margin = min(abs(v - W.THRESHOLD) for v in g["ratio"].values())
    R.check("SF5 every golden ratio at least 0.05 from 1.15, so exceeds reads one way", margin >= 0.05, "%.4f" % margin)

    # ---------------------------------------------------------------- the ladder on the pack
    for how, want in ROUTE_SETS.items():
        got = P.determination(how=how, frame=f)["losers"]
        R.check("L-%s files its named set" % how, got == sorted(want), str(got))
    for by, want in BY_SETS.items():
        got = P.determination(by=by, frame=f)["losers"]
        R.check("L-pooled-%s files its named set" % by, got == sorted(want), str(got))
    ha = P.determination(frame=f, mkw=dict(hosp_rates=True), by="all")
    R.check("L-hospital-rates-all adds Varden", ha["losers"] == sorted(["Varden Surgical Hospital"] + GOLD),
            str(ha["losers"]))
    hp = P.determination(frame=f.assign(**{"pass": np.where(f["cls"] == "ssn", "SSN", "NAME")}), by="pass",
                         mkw=dict(hosp_rates=True))
    R.check("L-hospital-rates-by-pass adds Varden", hp["losers"] == sorted(["Varden Surgical Hospital"] + GOLD),
            str(hp["losers"]))
    rf = P.determination(how="reflex", frame=f)
    R.check("L-reflex distinguishing margins", rf["ratio"]["Varden Surgical Hospital"] > 1.2
            and rf["ratio"]["Mercy Valley Medical Center"] < 1.1 and rf["ratio"]["Brannock Medical Center"] < 1.1,
            "%.3f %.3f %.3f" % (rf["ratio"]["Varden Surgical Hospital"], rf["ratio"]["Mercy Valley Medical Center"],
                                rf["ratio"]["Brannock Medical Center"]))
    # the readings that converge on the golden set (SF rows)
    f4 = f.assign(cls4=f["cls"] + np.where(f["comp"], "_c", "_s"))
    conv = {
        "SF8a four classes (number on file split by surname too)": P.determination(frame=f4, by="cls4"),
        "SF8b two-part surname read from a space only": P.determination(
            frame=f.assign(cls=np.where(f["cls"] == "ssn", "ssn", np.where(f["space"], "compound", "simple")))),
        "SF10 recorded deaths in hospital only as the found deaths": P.determination(frame=f, mkw=dict(ctrl="stay")),
        "SF11 each hospital's own class rates": P.determination(frame=f, mkw=dict(hosp_rates=True)),
        "SF12 the share linked alone as the divisor": P.determination(frame=f, mkw=dict(denom="r")),
        "SF13 contradicted links inside the uncontrolled admissions dropped": P.determination(how="refl_u", frame=f),
        "SF9a service lines alone as evidence of life": P.determination(alive_def="lines"),
        "SF9b evidence on day 30 itself counted": P.determination(alive_def="on30"),
        "SF14 links joined by member and date of death": P.determination(link_join="window"),
        "SF3 deaths on day 30 left out (29 days)": P.determination(window=29),
    }
    for k, d in conv.items():
        R.check(k + " files the golden set", d["losers"] == GOLD, str(d["losers"]))

    # ---------------------------------------------------------------- the analyst test on the clean truth
    tf = M.frame_from_truth(x)
    t_reflex = M.losers(M.ratios(tf, M.deaths(tf, how="reflex")))
    t_gold = M.losers(M.ratios(tf, M.deaths(tf)))
    R.check("R1 the easy route on the clean truth files the default path's set", t_reflex == REFLEX, str(t_reflex))
    R.check("R2 the decisive step alone on the clean truth files the golden set", t_gold == GOLD, str(t_gold))
    note = " ".join(pdf_text(os.path.join(target, W.FILES["linknote"])).split())
    R.check("SF2 the Office's note says its linkage misses deaths and links members to deaths not theirs",
            "some deaths of plan members are not found" in note and "linked to a death that is not theirs" in note)
    R.check("R3 every filed rule executed with the links as deaths files a wrong set",
            M.losers(M.ratios(tf, M.deaths(tf, how="naive"))) != GOLD)

    # ---------------------------------------------------------------- Layer 1 neutral on the call
    l1 = {
        "SF1a Carrisford's encounters left out": dict(plan_only=True),
        "SF1b voided bills and deleted encounters kept": dict(keep_voided=True),
        "SF6 transfers not joined (the receiving stay a later admission)": dict(join_transfers=False),
        "SF1c every eligible admission an index admission": dict(first_only=False),
        "SF1d deliveries kept": dict(drop_delivery=False),
        "SF1e discharges against advice kept": dict(drop_ama=False),
        "SF7 Mercy Valley's second NPI unmapped": dict(roster_drop=("1952407316",)),
    }
    for k, kw in l1.items():
        try:
            d = P.determination(**kw)
            R.check(k + " leaves the golden set", d["losers"] == GOLD, str(d["losers"]))
        except KeyError as e:          # an admission the scores file does not carry (the reading cannot be run)
            R.check(k + " cannot be scored, so it cannot move the call", True, str(e))

    # ---------------------------------------------------------------- the bonus figures and their device
    gold_b = {h["name"]: h["bonus"] for h in fig["hospitals"]}
    b = P.bonus()
    r100 = lambda v: int(round(v / 100.0)) * 100
    R.check("A1 golden bonus figures recompute", all(r100(b[h]) == gold_b[h] for h in W.CANDIDATES))
    off = {h: b[h] - round(b[h] / 100.0) * 100 for h in W.CANDIDATES}
    edge = min(50 - abs(v) for v in off.values())
    R.check("A2 every bonus figure at least 15 from its rounding edge", edge >= 15, "%.2f" % edge)
    R.check("A2b no bonus figure sits on a round number (each 12 or more from its cell's centre)",
            min(abs(v) for v in off.values()) >= 12, "%.2f" % min(abs(v) for v in off.values()))
    tot = sum(b.values())
    R.check("A3 the total rounds the same from the figures and from the payments, 25 or more from its edge",
            r100(tot) == sum(gold_b.values()) == fig["total_bonus"] and 50 - abs(tot - r100(tot)) >= 25, "%.2f" % tot)
    if "disputes_truth" not in x:
        __import__("pack").dispute_rows(x)
    truth = {h: 0.0 for h in W.CANDIDATES}
    seen = set()
    for bl in x["plan_bills"]:
        s = bl["stay"]
        if id(s) in seen or not (W.FY_START <= s["disch"] <= W.FY_END):
            continue
        seen.add(id(s))
        truth[s["hosp"]] += s["truth_paid"]
    for d in x["disputes_truth"]:
        if d["grp"] == "finance" and d["pay"] is not None and W.FY_START <= d["disch"] <= W.FY_END:
            truth[d["hosp"]] += d["award"]
    R.check("A4 the golden payments equal the clean truth (bills as last paid and Finance's awards)",
            all(abs(truth[h] - P.payments()[h]) < 0.01 for h in W.CANDIDATES))
    stops = {"A5 every paid award added, the reprocessed ones a second time (the cracker)": dict(awards="paid"),
             "A6 the dispute log left out": dict(awards="none"),
             "A7 every award decided for the hospital, paid or not": dict(awards="decided")}
    for k, kw in stops.items():
        v = P.bonus(**kw)
        hits = sum(r100(v[h]) == gold_b[h] for h in W.CANDIDATES) + (r100(sum(v.values())) == fig["total_bonus"])
        R.check(k + " misses all 15 figures", hits == 0, "%d hits" % hits)
    v = P.bonus(award_basis="payment")
    hits = sum(r100(v[h]) == gold_b[h] for h in W.CANDIDATES)
    R.check("A9 awards by the day they were paid misses at least 4 of the 14", hits <= 10, "%d hits" % hits)
    hazards = {"A8 the amount disputed": dict(awards="disputed"),
               "A10 a replacement standing in its original's place": dict(how="standard"),
               "A11 every line summed": dict(how="sum"), "A12 every replacement added": dict(how="net"),
               "A13 admission date as the basis": dict(basis="admission"),
               "A14 Carrisford's payments added": dict(with_cmg=True)}
    for k, kw in hazards.items():
        v = P.bonus(**kw)
        hits = sum(r100(v[h]) == gold_b[h] for h in W.CANDIDATES)
        R.check(k + " misses at least 10 of the 14", hits <= 4, "%d hits" % hits)
    v = P.bonus(npi_map=False)
    R.check("A15 Mercy Valley's second NPI unmapped misses Mercy Valley",
            r100(v["Mercy Valley Medical Center"]) != gold_b["Mercy Valley Medical Center"])
    # the device is fair: the records show which awards a bill already carries, to the cent and the day
    chains = {c["root"]: c for c in P.plan_chains()}
    after = {}
    cl = P.claims()
    for r in cl[cl["type_of_bill"] != "0111"].itertuples(index=False):
        after.setdefault(r.original_claim_number, []).append(B.d_iso(r.paid_date))
    carried, apart, near = 0, 0, 0
    for d in P.disputes():
        if d["Determination"] not in ("Overturned", "Partially overturned") or d["Payment Date"] is None:
            continue
        c = chains.get(str(d["Claim No."]))
        if c is None:
            continue
        a = float(d["Additional Payment"])
        if c["rep"] is not None and c["rep_date"] == d["Payment Date"]:
            change = c["rep"] - c["a0"] if c["rep_new"] else c["rep"]
            carried += abs(change - a) < 0.005
            near += abs(change - a) >= 0.005
        elif not any(t >= d["Determination Date"] for t in after.get(str(d["Claim No."]), [])):
            apart += 1
        else:
            near += 1
    R.check("A16 every paid award either is the replacement bill paid that day (its change the award to the cent) or "
            "has no bill after its determination", near == 0 and carried >= 150 and apart >= 150,
            "%d carried, %d apart, %d neither" % (carried, apart, near))
    c = P.claims()
    c = c.assign(amt=c["paid_amount"].astype(float), hosp=c["billing_npi"].map(P.roster()))
    rep = c[c["type_of_bill"] == "0117"]
    orig = c.set_index("claim_number")
    replaced = set(rep["original_claim_number"]) | set(c.loc[c["type_of_bill"] == "0118", "original_claim_number"])
    plain = c[(c["type_of_bill"] == "0111") & ~c["claim_number"].isin(replaced)]
    rate = plain.groupby(["hosp", "drg"])["amt"].agg(["min", "max"])
    R.check("A17 every uncorrected bill pays its hospital's one amount for its DRG",
            bool((rate["min"] == rate["max"]).all()))
    amt = rate["min"].to_dict()
    cps = rep[~rep["claim_number"].str.startswith("HC")]
    plat = rep[rep["claim_number"].str.startswith("HC")]
    cps_ok = sum(abs(r.amt - (amt[(r.hosp, r.drg)] - amt[(r.hosp, orig.loc[r.original_claim_number, "drg"])])) < 0.02
                 for r in cps.itertuples() if (r.hosp, r.drg) in amt
                 and (r.hosp, orig.loc[r.original_claim_number, "drg"]) in amt)
    plat_ok = sum(abs(r.amt - amt[(r.hosp, r.drg)]) < 0.02 for r in plat.itertuples() if (r.hosp, r.drg) in amt)
    R.check("A18 CPS replacements carry the change between the two DRGs' amounts (the filed convention)",
            cps_ok >= 0.95 * len(cps), "%d of %d" % (cps_ok, len(cps)))
    R.check("A19 platform replacements carry the full amount for the new DRG (the filed convention)",
            plat_ok >= 0.95 * len(plat), "%d of %d" % (plat_ok, len(plat)))
    R.check("A20 CPS voids carry the voided amount as a positive, platform voids as a negative",
            bool((c[(c["type_of_bill"] == "0118") & ~c["claim_number"].str.startswith("HC")]["amt"] > 0).all()) and
            bool((c[(c["type_of_bill"] == "0118") & c["claim_number"].str.startswith("HC")]["amt"] < 0).all()))
    notes = " ".join(pdf_text(os.path.join(target, W.FILES["claimnotes"])).split())
    R.check("A21 the extract notes file both systems' replacement and void conventions",
            "carries the change CPS made to the payment" in notes and "carries the stay's full payment" in notes
            and "positive amount on a CPS void and a negative amount on a platform void" in notes)
    # C1: no hospital's class short of its expected false links, so a floor on the estimate never binds
    Rr = M.rates(f, "cls")
    U = f[~f["known_death"] & ~f["alive_after"]]
    short = [(h, k) for (h, k), u in U.groupby(["hosp", "cls"]) if u["linked"].sum() < Rr[k][1] * len(u)]
    R.check("C1 every hospital's class carries at least its expected false links (floored and unfloored agree)",
            not short, str(short))
    d0 = P.determination(frame=f)
    d1 = M.deaths(f.assign(), how="golden")
    R.check("C1b the deaths with no floor equal the golden deaths", all(abs(d0["deaths"][h] - d1[h]) < 1e-9
                                                                        for h in W.CANDIDATES))
    enc = P.encounters()
    R.check("E1 every encounter was submitted by September 30, before the file was created on October 6",
            max(B.d_us2(v) for v in enc["submit"]) <= W.CLAIMS_THROUGH)
    # separation: the call never reads a paid amount
    P2 = B.Pack(target)
    P2._c["claims"] = P.claims().assign(paid_amount="0.00")
    R.check("S1 the call is unchanged with every paid amount blanked",
            P2.determination()["losers"] == GOLD and abs(P2.determination()["ratio"]["Mercy Valley Medical Center"] -
                                                         g["ratio"]["Mercy Valley Medical Center"]) < 1e-12)
    P3 = B.Pack(target)
    P3._c["pdr"] = []
    R.check("S3 the call is unchanged with the dispute log emptied", P3.determination()["losers"] == GOLD and
            all(abs(P3.determination()["ratio"][h] - g["ratio"][h]) < 1e-12 for h in W.CANDIDATES))
    src = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "bundle.py"), encoding="utf-8").read()
    R.check("S2 the solution never reads the distractors",
            all('"%s"' % k not in src for k in W.DISTRACTORS))

    # ---------------------------------------------------------------- the deliverable and the gates
    wb = openpyxl.load_workbook(os.path.join(golden_dir, W.GOLDEN["workbook"]), data_only=True)
    ws = wb["Determination"]
    rows = {ws.cell(row=i, column=1).value: (ws.cell(row=i, column=8).value, ws.cell(row=i, column=10).value)
            for i in range(5, 19)}
    R.check("D1 the workbook's verdicts and bonus figures equal figures.json",
            all(rows[h["name"]] == ("Loses" if h["verdict"] == "loses" else "Keeps", h["bonus"])
                for h in fig["hospitals"]))
    files = sorted(os.listdir(target))
    R.check("G1 ten or more files in three or more formats", len(files) >= 10 and
            len({fn.rsplit(".", 1)[1] for fn in files}) >= 3)
    rows_big = sum(1 for _ in open(os.path.join(target, W.FILES["activity"]), encoding="utf-8")) - 1
    R.check("G2 a file of 25,000 rows or more", rows_big >= 25000, str(rows_big))
    meta = json.load(open(meta_path, encoding="utf-8"))
    R.check("G3 two distractors named in metadata.json", len(meta["distractor_files"]) == 2)
    text = ""
    for fn in files:
        p = os.path.join(target, fn)
        if fn.endswith(".pdf"):
            text += pdf_text(p).lower()
        elif fn.endswith(".xlsx"):
            wbk = openpyxl.load_workbook(p, read_only=True)
            for wsx in wbk.worksheets:
                if wsx.title == "Scores":
                    continue
                for r in wsx.iter_rows(values_only=True):
                    text += " ".join(str(v) for v in r if v is not None).lower() + "\n"
        else:
            with open(p, encoding="utf-8") as fh:
                text += fh.readline().lower()
    found = [w for w in DECISIVE_WORDS if w in text]
    R.check("R3c no document names the decisive step", not found, str(found))
    R.check("T1 no en or em dash in the pack's documents", not re.search("[–—]", text))
    print("checks: %d run, %d failed" % (R.n, len(R.bad)))
    return not R.bad
