"""The pack's documents: SBA Lending's allowance policy and the bank's SBA loan authorization standard terms (PDF), the
README with its field notes, the export manifest and the workbooks (guarantee register, the agent's Economic Aid Act
layout). The documents' text is kept here as data so that checks.py can quote and execute it."""
from __future__ import annotations

import datetime as dt
import re

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import letter                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import inch                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer                    # noqa: E402

import common as K     # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/"
for _n, _f in (("Carlito", "crosextra/Carlito-Regular.ttf"), ("Carlito-Bold", "crosextra/Carlito-Bold.ttf"),
               ("LibSerif", "liberation/LiberationSerif-Regular.ttf"),
               ("LibSerif-Bold", "liberation/LiberationSerif-Bold.ttf")):
    pdfmetrics.registerFont(TTFont(_n, FONT_DIR + _f))
INK = colors.HexColor("#1b2430")
GREY = colors.HexColor("#5b6470")
NAVY = colors.HexColor("#163b5c")       # Fenmore's navy

BODY = ParagraphStyle("body", fontName="Carlito", fontSize=10.4, leading=13.8, textColor=INK, spaceAfter=5)
H1 = ParagraphStyle("h1", fontName="Carlito-Bold", fontSize=15.0, leading=19, textColor=NAVY, spaceAfter=2)
H1S = ParagraphStyle("h1s", fontName="Carlito", fontSize=9.2, leading=12.0, textColor=GREY, spaceAfter=10)
H2 = ParagraphStyle("h2", fontName="Carlito-Bold", fontSize=11.2, leading=14.5, textColor=NAVY, spaceBefore=8, spaceAfter=3)
LG_BODY = ParagraphStyle("lgb", fontName="LibSerif", fontSize=10.0, leading=13.2, textColor=INK, spaceAfter=4)
LG_H = ParagraphStyle("lgh", fontName="LibSerif-Bold", fontSize=10.4, leading=13.6, textColor=INK, spaceBefore=6,
                      spaceAfter=2)
LG_T = ParagraphStyle("lgt", fontName="LibSerif-Bold", fontSize=13.0, leading=17, textColor=INK, spaceAfter=3,
                      alignment=1)
LG_S = ParagraphStyle("lgs", fontName="LibSerif", fontSize=9.6, leading=12.6, textColor=INK, spaceAfter=10, alignment=1)


def _doc(path, title, author, margins=1.0):
    return SimpleDocTemplate(path, pagesize=letter, leftMargin=margins * inch, rightMargin=margins * inch,
                             topMargin=0.85 * inch, bottomMargin=0.85 * inch, title=title, author=author, subject="",
                             creator="", producer="")


def _footer(left, right, font="Carlito"):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 7.8)
        canvas.setFillColor(GREY)
        canvas.drawString(1.0 * inch, 0.5 * inch, left)
        canvas.drawRightString(letter[0] - 1.0 * inch, 0.5 * inch, "%s%s" % (right, doc.page))
        canvas.restoreState()
    return draw


# ---------------------------------------------------------------------------------------------------- the policy
F_POLICY = "allowance_policy_sba_7a.pdf"
POLICY = [
    ("1", "Scope", "scope", [
        "1.1 This policy sets the allowance for credit losses on the unguaranteed portions of the SBA 7(a) loans SBA "
        "Lending holds. Guaranteed portions sold on the secondary market are not on the bank's books; guaranteed portions "
        "the bank holds are covered by SBA's guaranty and carry no allowance.",
        "1.2 The relief-era pool is the 7(a) loans approved by SBA in its fiscal years 2019 to 2021, from 1 October 2018 "
        "to 30 September 2021. Its allowance is estimated and reported separately from the allowance on the other "
        "vintages."]),
    ("2", "The estimate", "estimate", [
        "2.1 The allowance for the pool at a reserve date is the unguaranteed principal the bank expects to charge off "
        "over each loan's remaining contractual term, adjusted for expected prepayments, on the pool's loans on the books "
        "at that date, loans in default not yet charged off included. Expected recoveries are not included.",
        "2.2 A loan is in default at the first month-end at which it is 90 or more days past due or has been accelerated.",
        "2.3 The unguaranteed principal of a loan in default is charged off in full when SBA has purchased the "
        "guaranteed portion, unless the loan has been brought current or paid; amounts collected afterwards are "
        "recoveries."]),
    ("3", "Exposure", "exposure", [
        "3.1 A loan's exposure at a month-end after the reserve date is its unguaranteed share of the principal scheduled "
        "to be outstanding at that month-end, on its amortization schedule at the note rate in effect at the reserve "
        "date."]),
    ("4", "Experience and conditions", "experience", [
        "4.1 The estimate draws on the bank's own experience with its SBA 7(a) loans from October 2009, when the division "
        "opened. Management's reasonable and supportable forecast of conditions over the pool's remaining lives does not "
        "differ from that experience, and no adjustment is made for forecast conditions."]),
    ("5", "Run-off", "runoff", [
        "5.1 The estimate covers the loans on the books at the reserve date. Loans the bank may make after that date are "
        "not projected."]),
    ("6", "Governance", "governance", [
        "6.1 The chief credit officer, SBA Lending, owns this policy and reviews it each year. The controller books the "
        "allowance at each 30 September, the fiscal year-end, and the board approves it at its October meeting.",
        "6.2 Changes to this policy are approved by the board's risk committee."]),
]


def policy_sentences(key=None):
    out = []
    for _, _, k, paras in POLICY:
        if key is None or k == key:
            out += [s for p in paras for s in re.split(r"(?<=\.)\s+(?=[A-Z0-9])", p)]
    return out


def write_policy(path):
    s = [Paragraph("Fenmore Savings Bank &middot; SBA Lending &middot; Credit policy CP-SBA-07, version 4", H1S),
         Paragraph("SBA 7(a) unguaranteed portions: allowance for credit losses", H1),
         Paragraph("Approved by the Risk Committee of the Board on 23 July 2024; effective for reserve dates from "
                   "30 September 2024. Owner: Richard Fields, Chief Credit Officer, SBA Lending.", H1S)]
    for num, head, _, paras in POLICY:
        s.append(Paragraph("%s&nbsp;&nbsp;%s" % (num, head), H2))
        for p in paras:
            s.append(Paragraph(p.replace("&", "&amp;"), BODY))
    doc = _doc(path, "CP-SBA-07 SBA 7(a) unguaranteed portions: allowance for credit losses", K.PEOPLE["cco"])
    f = _footer("Fenmore Savings Bank  |  CP-SBA-07 v4  |  Internal", "Page ")
    doc.build(s, onFirstPage=f, onLaterPages=f)


# ---------------------------------------------------------------------------------------------------- the terms
F_TERMS = "sba_loan_authorization_standard_terms.pdf"
TERMS = [
    ("1", "Repayment and interest", "repayment", [
        "1.1 Borrower shall repay the Loan in monthly instalments of principal and interest, beginning on the first "
        "payment date in the Note, in amounts that repay the Loan by its maturity date.",
        "1.2 A variable rate is the Wall Street Journal Prime Rate in effect on the first business day of the month in "
        "which SBA received the loan application, or on the adjustment date, plus the spread in the Note. The rate "
        "adjusts on the first day of January, April, July and October, and the instalment is recalculated at each "
        "adjustment so that the Loan is repaid by its maturity date.",
        "1.3 Borrower may prepay the Loan in whole or in part at any time without premium."]),
    ("2", "Collateral", "collateral", [
        "2.1 The Loan is secured by a first lien on all business assets of Borrower, now owned or later acquired, and by "
        "any other collateral named in Part A.",
        "2.2 Lender may require an assignment of life insurance on the principal operator of the business."]),
    ("3", "Guarantees", "guarantees", [
        "3.1 Each individual and each entity owning 20 per cent or more of Borrower shall guarantee the Loan, unlimited, "
        "jointly and severally. Lender may require a limited guarantee of an owner of less than 20 per cent.",
        "3.2 Each guarantee is a guarantee of payment and continues until the Loan is paid in full."]),
    ("4", "Cross-collateralization and cross-default", "crossdefault", [
        "4.1 Every loan Lender has made or makes to Borrower, to any Affiliate of Borrower, or to any person or entity "
        "whose obligations to Lender are guaranteed by a Guarantor of this Loan is cross-collateralized with this Loan.",
        "4.2 A default under any of those loans is a default under this Loan."]),
    ("5", "Events of default", "default", [
        "5.1 Borrower is in default under the Note if Borrower or a Guarantor fails to make a payment when due; fails to "
        "comply with any term of the Loan documents; ceases operations, sells substantially all its assets or is "
        "liquidated; becomes the subject of a bankruptcy or insolvency proceeding; makes a materially false statement to "
        "Lender or SBA; or if a default occurs under section 4."]),
    ("6", "Lender's rights", "rights", [
        "6.1 On a default Lender may, without notice, declare the whole Loan immediately due and payable, collect from "
        "any Guarantor, and take possession of and sell the collateral."]),
    ("7", "Reporting and insurance", "reporting", [
        "7.1 Borrower shall give Lender annual financial statements within 120 days of each fiscal year-end and copies of "
        "its federal tax returns when filed.",
        "7.2 Borrower shall keep the collateral insured against loss for its full insurable value, Lender named as loss "
        "payee."]),
    ("8", "SBA requirements", "sba", [
        "8.1 The Loan is made under SBA's 7(a) program and is subject to SBA's Loan Program Requirements. Borrower shall use "
        "the proceeds only for the purposes in Part A and shall sign SBA Form 1050 at each disbursement.",
        "8.2 Borrower may not assume, transfer or sell the business financed without Lender's consent and, where SBA "
        "requires it, SBA's."]),
]


def terms_sentences(key=None):
    out = []
    for _, _, k, paras in TERMS:
        if key is None or k == key:
            out += [s for p in paras for s in re.split(r"(?<=\.)\s+(?=[A-Z0-9])", p)]
    return out


def write_terms(path):
    s = [Paragraph("Fenmore Savings Bank", LG_T),
         Paragraph("SBA LOAN AUTHORIZATION, PART B: STANDARD TERMS<br/>Form SBL-AUTH-B, revised March 2018. Part B applies "
                   "to every SBA 7(a) loan Lender authorizes under the Preferred Lenders Program; Part A states each "
                   "loan's amount, term, rate, collateral and guarantors.", LG_S)]
    for num, head, _, paras in TERMS:
        s.append(Paragraph("%s. %s" % (num, head), LG_H))
        for p in paras:
            s.append(Paragraph(p.replace("&", "&amp;"), LG_BODY))
    doc = _doc(path, "SBA Loan Authorization Part B Standard Terms", "Fenmore Savings Bank SBA Lending", margins=1.1)
    f = _footer("Fenmore Savings Bank  |  SBL-AUTH-B (03/2018)", "Page ", font="LibSerif")
    doc.build(s, onFirstPage=f, onLaterPages=f)


# ---------------------------------------------------------------------------------------------------- workbooks
def write_register(path, cols, rows):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Guaranties"
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r[c] for c in cols])
    for col, w in zip("ABCDEFGHIJ", (14, 44, 12, 36, 11, 12, 12, 11, 12, 12)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=9, max_col=10):
        for c in row:
            if c.value is not None:
                c.number_format = "mm/dd/yyyy"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:J%d" % (len(rows) + 1)
    wb.save(path)


def write_eaa(path, cols, rows):
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1112 EAA Detail"
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r[c] for c in cols])
    for col, w in zip("ABCDEF", (12, 18, 15, 44, 12, 22)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        for c in row:
            c.number_format = "m/d/yyyy"
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        for c in row:
            c.number_format = "m/d/yyyy"
    ws.freeze_panes = "A2"
    wb.save(path)


# ---------------------------------------------------------------------------------------------------- README
F_README = "README.txt"


def readme_text(files):
    L = []
    L.append("SBA LENDING - DATA FOR THE 30 SEPTEMBER 2026 ALLOWANCE")
    L.append("Fenmore Savings Bank. Prepared by J. Nicholson (SBA servicing) and T. Williams (loan operations), 5 October 2026.")
    L.append("")
    L.append("FILES")
    for f, d in files:
        L.append("%s" % f)
        L.append("    %s" % d)
    L.append("")
    L.append("FIELD NOTES")
    notes = [
        ("servicing_history_2017-10_2026-09.parquet", [
            "One row per loan per month-end from October 2017, the month the current servicing system went live, to "
            "September 2026. A loan's rows run from the month-end of its first disbursement to the month it closes "
            "(charged off, paid in full or matured).",
            "note_no: the loan's note number on the current system. status: see servicing_status_codes.csv. dpd: days "
            "past due at the month-end, blank on liquidation and closing rows. principal_bal: principal outstanding at "
            "the month-end, dollars; 0.00 on the closing row. note_rate: the interest rate in effect at the month-end, "
            "per cent a year."]),
        ("legacy_loan_history_200910_201710.parquet", [
            "One row per loan per month-end from October 2009 to October 2017 from the legacy core. October 2017 was run "
            "on both systems and is in both extracts.",
            "LOAN_NO: the legacy loan number. PERIOD: year and month of the month-end (YYYYMM). STAT_CD: see "
            "servicing_status_codes.csv. DAYS_DLQ: days past due at the month-end (0 on liquidation and closing rows). "
            "PRIN_BAL_CENTS: principal outstanding, cents. INT_RT: interest rate as a decimal fraction."]),
        ("sba_loan_master.csv", [
            "One row per SBA 7(a) loan approved from October 2009 to September 2026 and disbursed by 30 September 2026. "
            "A loan carries its legacy number if it was booked before the conversion and its note number if it was on "
            "the current system. approval_date: SBA's approval date. guaranty_pct: SBA's guaranteed share, per cent. "
            "rate_type: V variable (prime plus spread, adjusted quarterly), F fixed. initial_rate: per cent a year."]),
        ("guarantee_register.xlsx", [
            "One row per guaranty: each guarantor of each loan, as named on the guaranty. Party No: the documentation "
            "system's customer number for the guarantor. TIN (last 4): the last four digits of the guarantor's taxpayer "
            "identification number (SSN or EIN). Released: the date the guaranty was released on payment of the loan."]),
        ("fta_1112_detail_cares.csv and fta_1112_detail_eaa.xlsx", [
            "SBA's fiscal transfer agent's remittance detail for Section 1112 payments, as received: the CARES Act "
            "layout and the Economic Aid Act layout. Every vintage of the 7(a) book. Loans are keyed by SBA loan number."]),
        ("sba_settlement_account_credits.csv", [
            "ACH and wire credits to the SBA settlement account (DDA ending 4471) from SBA's fiscal transfer agent and "
            "SBA's Denver Finance Center, October 2019 to September 2026, with the originator's addenda."]),
        ("special_assets_ledger_fy2020_fy2026.csv", [
            "Special Assets' charge-off and recovery entries on SBA 7(a) loans posted from 1 October 2019 to 30 September "
            "2026. amount: dollars; charge-offs positive, receipts and reversals negative."]),
        ("payments_to_sba.csv", [
            "Payments from the bank to SBA, October 2019 to September 2026, by wire or ACH."]),
    ]
    for f, ps in notes:
        L.append(f)
        for p in ps:
            L += ["    " + x for x in _wrap(p, 104)]
    L.append("")
    L.append("Fiscal years run from 1 October to 30 September (FY2026: 1 October 2025 to 30 September 2026); SBA's fiscal "
             "year is the same.")
    return "\r\n".join(L) + "\r\n"


def _wrap(text, width):
    words, lines, cur = text.split(), [], ""
    for w in words:
        if len(cur) + len(w) + 1 > width:
            lines.append(cur)
            cur = w
        else:
            cur = (cur + " " + w).strip()
    if cur:
        lines.append(cur)
    return lines
