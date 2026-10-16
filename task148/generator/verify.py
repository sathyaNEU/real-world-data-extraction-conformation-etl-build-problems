"""Independent verifier for task148: reads only the shipped bundle (target/) and the goldens, on its own code path, and
recomputes the allowance, its split, the unguaranteed principal, the Section 1112 table and the net charge-off table.

    python3 task148/generator/verify.py [target] [golden]

Groups are the connected components of the loans-by-guarantor graph (scipy); a loan's default probability in a month
is its survival times the chance that the first failure among its group's performing loans falls in that month, built
from cumulative products rather than per-step hazards; SBA's share of a receipt is read from the ratio of SBA's
remittance to the ledger amount, not from the module's date."""
from __future__ import annotations

import os
import re
import sys

import numpy as np
import pandas as pd

M0 = 2009 * 12 + 9          # month index origin: October 2009
T = 203                     # September 2026


def mi(ts):
    ts = pd.to_datetime(ts)
    return (ts.dt.year * 12 + ts.dt.month - 1 - M0).values


def load(target):
    import openpyxl
    import pyarrow.parquet as pq
    m = pd.read_csv(os.path.join(target, "sba_loan_master.csv"), dtype=str)
    key = {}
    for col in ("legacy_loan_no", "note_no"):
        s = m[[col, "sba_loan_no"]].dropna()
        key.update(dict(zip(s[col], s.sba_loan_no)))
    cur = pq.read_table(os.path.join(target, "servicing_history_2017-10_2026-09.parquet")).to_pandas()
    leg = pq.read_table(os.path.join(target, "legacy_loan_history_200910_201710.parquet")).to_pandas()
    codes = pd.read_csv(os.path.join(target, "servicing_status_codes.csv"))
    desc = dict(zip(codes.code, codes.description))
    cur_map = {c: desc[c] for c in cur.status.unique()}
    leg_map = {c: desc[c] for c in leg.STAT_CD.unique()}
    a = pd.DataFrame({"loan": cur.note_no.map(key), "m": mi(cur.period_end), "st": cur.status.map(cur_map),
                      "bal": cur.principal_bal, "rate": cur.note_rate})
    per = leg.PERIOD.astype(int)
    b = pd.DataFrame({"loan": leg.LOAN_NO.map(key), "m": (per // 100) * 12 + per % 100 - 1 - M0,
                      "st": leg.STAT_CD.map(leg_map), "bal": leg.PRIN_BAL_CENTS / 100.0, "rate": leg.INT_RT * 100})
    x = pd.concat([b, a]).drop_duplicates(["loan", "m"], keep="last").sort_values(["loan", "m"]).reset_index(drop=True)
    wb = openpyxl.load_workbook(os.path.join(target, "guarantee_register.xlsx"), read_only=True)
    rows = list(wb["Guaranties"].iter_rows(values_only=True))
    reg = pd.DataFrame(rows[1:], columns=rows[0])
    return m, x, reg


PERFORMING = {"Current", "30 to 59 days past due", "60 to 89 days past due"}
DEFAULTED = {"90 or more days past due", "Accelerated (loan called due)"}


def allowance(target):
    from scipy.sparse import coo_matrix
    from scipy.sparse.csgraph import connected_components
    m, x, reg = load(target)
    loans = np.array(sorted(m.sba_loan_no))
    li = {s: k for k, s in enumerate(loans)}
    parties = {p: k for k, p in enumerate(sorted(reg["Party No"].unique()))}
    rr = reg["SBA Loan No"].map(li).values
    cc = reg["Party No"].map(parties).values + len(loans)
    n = len(loans) + len(parties)
    g = coo_matrix((np.ones(len(rr)), (rr, cc)), shape=(n, n))
    _, comp = connected_components(g, directed=False)
    grp = comp[:len(loans)]
    x["L"] = x.loan.map(li)
    x["grp"] = grp[x.L.values]
    first = x.groupby("L").m.transform("min")
    x["age"] = x.m - first
    prev = x.groupby("L").st.shift(1)
    x["atrisk"] = prev.isin(PERFORMING) & (x.age >= 1)
    x["isdef"] = x.atrisk & x.st.isin(DEFAULTED)
    x["ispo"] = x.atrisk & (x.st == "Paid in full")
    r = x[x.atrisk].copy()
    r["size"] = r.groupby(["grp", "m"]).L.transform("size")
    bands = 25
    r["b"] = np.minimum((r.age - 1) // 6, bands - 1)

    def table(d, col):
        e = np.bincount(d.b, weights=d[col].astype(float), minlength=bands)
        nn = np.bincount(d.b, minlength=bands).astype(float)
        out = np.full(bands, np.nan)
        out[nn > 0] = e[nn > 0] / nn[nn > 0]
        return pd.Series(out).ffill().fillna(0).values
    hz = table(r[r["size"] == 1], "isdef")
    pz = table(r, "ispo")
    # charge-off share: the first status after each default that is not liquidation
    st = x.st.values
    Lc = x.L.values
    ends = np.r_[np.where(np.diff(Lc) != 0)[0], len(Lc) - 1]
    end_of = np.repeat(ends, np.diff(np.r_[-1, ends]))
    res = []
    for p in np.where(x.isdef.values)[0]:
        q = p + 1
        while q <= end_of[p] and st[q] == "In liquidation":
            q += 1
        if q <= end_of[p]:
            res.append(st[q] == "Charged off")
    share = float(np.mean(res))
    last = x.groupby("L").tail(1)
    book = last[(last.m == T) & ~last.st.isin(["Charged off", "Paid in full", "Matured, paid at maturity"])]
    ms = m.set_index("sba_loan_no")
    appr = pd.to_datetime(ms.approval_date)
    fy = pd.Series(np.where(appr.dt.month >= 10, appr.dt.year + 1, appr.dt.year), index=ms.index)
    mat = pd.Series(mi(ms.maturity_date), index=ms.index)
    ugs = 1 - ms.guaranty_pct.astype(float) / 100.0
    book = book.assign(fy=fy.loc[loans[book.L.values]].values, ug=ugs.loc[loans[book.L.values]].values,
                       rem=np.maximum(mat.loc[loans[book.L.values]].values - T, 0), age_now=T - first.groupby(x.L).first().loc[book.L].values)
    pool = book[(book.fy >= 2019) & (book.fy <= 2021)]
    perf = book[book.st.isin(PERFORMING)]
    total = {}
    for _, row in pool[~pool.st.isin(PERFORMING)].iterrows():
        total[row.L] = share * row.bal * row.ug
    for gid, d in perf[perf.grp.isin(set(pool.grp))].groupby("grp"):
        R = int(d.rem.max())
        if R <= 0:
            continue
        s = np.arange(1, R + 1)
        ages = d.age_now.values[:, None] + s[None, :]
        live = s[None, :] <= d.rem.values[:, None]
        bidx = np.minimum((ages - 1) // 6, bands - 1)
        h = np.where(live, hz[bidx], 0.0)
        p = np.where(live & (s[None, :] < d.rem.values[:, None]), pz[bidx], 0.0)
        surv = np.cumprod(np.c_[np.ones(len(d)), 1 - h - p], axis=1)          # alive and unfailed at the start of s
        failed_by = np.cumsum(surv[:, :-1] * h, axis=1)                       # failed by the end of s
        untrig = np.c_[np.ones(len(d)), 1 - failed_by]                        # not failed before s (index s-1), through s (index s)
        for j, row in enumerate(d.itertuples()):
            if not ((row.fy >= 2019) and (row.fy <= 2021)):
                continue
            rj = int(row.rem)
            if rj <= 0:
                total[row.L] = 0.0
                continue
            others = [k for k in range(len(d)) if k != j]
            A0 = np.prod(untrig[others][:, :rj], axis=0) if others else np.ones(rj)
            A1 = np.prod(untrig[others][:, 1:rj + 1], axis=0) if others else np.ones(rj)
            pj = surv[j, :rj] * (A0 - (1 - h[j, :rj]) * A1)
            rate = row.rate / 1200.0
            k = np.arange(1, rj + 1)
            ead = row.bal * ((1 + rate) ** rj - (1 + rate) ** k) / ((1 + rate) ** rj - 1)
            total[row.L] = float((pj * share * ead).sum() * row.ug)
    tot = pd.Series(total)
    out = {"allowance": float(tot.sum()), "share": share, "pool": len(pool)}
    out["split"] = {f: float(tot[tot.index.isin(pool.L[pool.fy == f])].sum()) for f in (2019, 2020, 2021)}
    out["ugp"] = {f: float((pool.bal * pool.ug)[pool.fy == f].sum()) for f in (2019, 2020, 2021)}
    return out


def section_1112(target):
    import openpyxl
    m = pd.read_csv(os.path.join(target, "sba_loan_master.csv"), dtype=str)
    appr = pd.to_datetime(m.approval_date)
    fy = dict(zip(m.sba_loan_no, np.where(appr.dt.month >= 10, appr.dt.year + 1, appr.dt.year)))
    c = pd.read_csv(os.path.join(target, "fta_1112_detail_cares.csv"), dtype=str)
    c = pd.DataFrame({"trace": c.trace_no, "loan": c.sba_loan_no, "due": pd.to_datetime(c.installment_due),
                      "amt": c.pmt_amount.astype(float)})
    wb = openpyxl.load_workbook(os.path.join(target, "fta_1112_detail_eaa.xlsx"), read_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    e = pd.DataFrame(rows[1:], columns=rows[0])
    e = pd.DataFrame({"trace": e["ACH Trace"].astype(str), "loan": e["SBA Loan Number"].str.replace("-", ""),
                      "due": pd.to_datetime(e["Due Date"]), "amt": e["Payment Amount (cents)"].astype(float) / 100})
    x = pd.concat([c, e])
    x = x[x.loan.map(fy).between(2019, 2021)]
    # one disbursement per trace and loan: its amount once, spread over the instalments it lists
    disb = x.groupby(["trace", "loan"]).agg(amt=("amt", "first"), n=("amt", "size"))
    x = x.join(disb, on=["trace", "loan"], rsuffix="_d")
    x["each"] = x.amt_d / x.n
    x = x[(x.due >= "2020-04-01") & (x.due <= "2021-12-31")]
    fyq = (x.due.dt.year + (x.due.dt.month >= 10)).astype(str) + " Q" + (((x.due.dt.month - 10) % 12) // 3 + 1).astype(str)
    return {"FY" + k: float(v) for k, v in x.groupby(fyq.values).each.sum().items()}


def net_charge_offs(target):
    m = pd.read_csv(os.path.join(target, "sba_loan_master.csv"), dtype=str)
    appr = pd.to_datetime(m.approval_date)
    fy = dict(zip(m.sba_loan_no, np.where(appr.dt.month >= 10, appr.dt.year + 1, appr.dt.year)))
    gty = dict(zip(m.sba_loan_no, m.guaranty_pct.astype(float) / 100))
    num = {}
    for col in ("note_no", "legacy_loan_no"):
        s = m[[col, "sba_loan_no"]].dropna()
        num.update(dict(zip(s[col], s.sba_loan_no)))
    L = pd.read_csv(os.path.join(target, "special_assets_ledger_fy2020_fy2026.csv"), dtype={"note_no": str})
    L["loan"] = L.note_no.map(num)
    L["date"] = pd.to_datetime(L.posted)
    L = L[L.loan.map(fy).between(2019, 2021)]
    pay = pd.read_csv(os.path.join(target, "payments_to_sba.csv"), dtype={"sba_loan_no": str})
    pay = pay[pay.purpose == "SBA SHARE OF RECOVERY"]
    pay["date"] = pd.to_datetime(pay.paid)
    cr = pd.read_csv(os.path.join(target, "sba_settlement_account_credits.csv"), dtype=str)
    cr = cr[cr.addenda.str.contains("TOP OFFSET", na=False)]
    cr["loan"] = cr.addenda.str.extract(r"LN (\d+)")[0]
    cr["date"] = pd.to_datetime(cr.settlement_date)
    out = {}

    def put(d, loan, v):
        f = d.year + (1 if d.month >= 10 else 0)
        if 2023 <= f <= 2026:
            out[(f, int(fy[loan]))] = out.get((f, int(fy[loan])), 0.0) + v
    for r in L[L.code.isin(["CO", "REV"])].itertuples():
        put(r.date, r.loan, r.amount)
    for loan, d in L[L.code == "RCV"].groupby("loan"):
        p = pay[pay.sba_loan_no == loan].sort_values("date")
        used = set()
        for r in d.sort_values("date").itertuples():
            amt = -r.amount
            cand = p[(p.date >= r.date) & (p.date <= r.date + pd.Timedelta(days=40)) & ~p.index.isin(used)]
            kept = amt
            if len(cand):
                k = cand.index[0]
                share = float(p.at[k, "amount"])
                used.add(k)
                gross_ratio = gty[loan]                     # SBA's share of a receipt recorded gross
                if abs(share / amt - gross_ratio) < 0.01:
                    kept = amt - share
            put(r.date, r.loan, -kept)
    for r in cr.itertuples():
        if r.loan in fy and 2019 <= fy[r.loan] <= 2021:
            put(r.date, r.loan, -float(r.amount))
    return out


def golden_values(golden):
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(golden, "sba_allowance_fy2026.xlsx"), data_only=True)
    a = wb["Allowance"]
    rows = list(a.iter_rows(min_row=5, max_row=8, values_only=True))
    g = {"split": {int(r[0][2:]): r[4] * 1e6 for r in rows[:3]}, "allowance": rows[3][4] * 1e6,
         "ugp": {int(r[0][2:]): r[3] for r in rows[:3]}}
    s = wb["Section 1112"]
    g["ask_a"] = {r[0]: r[2] for r in s.iter_rows(min_row=4, values_only=True) if r[0] and r[0].startswith("FY")}
    n = wb["Net charge-offs"]
    g["ask_b"] = {}
    for r in n.iter_rows(min_row=4, values_only=True):
        if r[0] and str(r[0]).startswith("FY"):
            for k, v in zip((2019, 2020, 2021), r[1:4]):
                g["ask_b"][(int(r[0][2:]), k)] = v
    return g


def rnd(x, step):
    q = x / step
    return int(np.floor(q + 0.5)) * step if q >= 0 else -int(np.floor(-q + 0.5)) * step


def main(target, golden):
    v = allowance(target)
    a = section_1112(target)
    b = net_charge_offs(target)
    g = golden_values(golden)
    fails = []

    def eq(label, got, want):
        ok = got == want
        print("%-6s %-46s %s %s" % ("ok" if ok else "FAIL", label, got, "" if ok else "(golden %s)" % want))
        if not ok:
            fails.append(label)
    eq("allowance", rnd(v["allowance"], 1e6), g["allowance"])
    for f in (2019, 2020, 2021):
        eq("allowance FY%d" % f, rnd(v["split"][f], 5e5), g["split"][f])
        eq("unguaranteed principal FY%d" % f, rnd(v["ugp"][f], 1000), g["ugp"][f])
    for k, val in a.items():
        eq("Section 1112 %s" % k, rnd(val, 1000), g["ask_a"][k])
    for k in sorted(g["ask_b"]):
        eq("net charge-offs FY%d, FY%d approvals" % k, rnd(b.get(k, 0.0), 1000), g["ask_b"][k])
    print("verify: allowance %.2f (charge-off share %.4f, %d pool loans); %d figures, %d failed" % (
        v["allowance"], v["share"], v["pool"], 7 + len(a) + len(g["ask_b"]), len(fails)))
    if fails:
        raise SystemExit("verify failed: " + ", ".join(fails))
    return v, a, b


if __name__ == "__main__":
    here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    t = sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "target")
    gd = sys.argv[2] if len(sys.argv) > 2 else os.path.join(here, "golden")
    main(t, gd)
