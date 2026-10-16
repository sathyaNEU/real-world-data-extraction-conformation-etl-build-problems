"""The Tri-County Regional 911 Board's hourly volume forecast workbook for fiscal 2027 and 2028: one long table per
year for the region's four PSAPs by month, day of the week and hour beginning, the average 911 calls presented per
clock hour, with the part that is the Halvorsen move. The agency's Friday and Saturday late-shift cells for fiscal
2028 are centred so the golden's compliance sits on its whole per cent on the parameters the ACD history yields."""
from __future__ import annotations

import datetime as dt

import numpy as np
from scipy.optimize import brentq

import common as K
import levels as LV
import model as M
import world as WD

PSAPS = [("Tavish County ECA", 1.0), ("Wexley County ECC", 0.83), ("Garrick County 911", 0.58),
         ("Hollin Creek PD", 0.13)]
HALV_ALL = 3.48          # Halvorsen's calls per hour across the week, fiscal 2028
WEEK_MEAN = float(np.mean(K.WEEKDAY))
PUBLISHED = dt.date(2026, 9, 15)
READ_ME = [
    "Tri-County Regional 911 Board",
    "Hourly 911 volume forecast, fiscal years 2027 and 2028",
    "Published 15 September 2026 (planning, William Braun)",
    "",
    "Each row is a PSAP, a month, a day of the week and an hour beginning (local time). The forecast is the average "
    "number of 911 calls presented to the PSAP's queue in that clock hour on that day of the week in that month: "
    "calls transferred in from other PSAPs are counted, administrative lines and test calls are not.",
    "Fiscal 2028 includes Halvorsen Township's calls, which move from Wexley County Emergency Communications to the "
    "Tavish County Emergency Communications Agency on 1 July 2027, location-based routing sending the township's "
    "wireless calls to Tavish County from the same date. The Halvorsen column is the part of each figure that is the "
    "move (negative for Wexley County).",
    "The forecast is built from each PSAP's call history, the region's growth and the Halvorsen move. It forecasts "
    "calls presented, not answer times or staffing.",
    "Fiscal years run from 1 July to 30 June.",
]


def _rep_date(y, m, dow):
    d = dt.date(y, m, 1)
    while d.weekday() != dow:
        d += dt.timedelta(days=1)
    return d + dt.timedelta(days=7)


def base_cell(y, m, dow, h):
    """The agency's own expected calls per hour in a cell (no Halvorsen), from the history's model run forward."""
    d = _rep_date(y, m, dow)
    loc = dt.datetime(d.year, d.month, d.day, h, 0, tzinfo=K.TZ)
    return WD.lam_hist(loc)


def halv_cell(m, dow, h):
    if (h >= 22 and dow in (4, 5)) or (h < 2 and (dow - 1) % 7 in (4, 5)):
        return LV.HALV_LATE[m] * LV.cell_factor(dow, h)
    return HALV_ALL * K.WEEKDAY[h] / WEEK_MEAN * K.SEAS[m]


def is_late_cell(dow, h):
    return (h >= 22 and dow in (4, 5)) or (h < 2 and (dow - 1) % 7 in (4, 5))


def centre_late(busy, pat):
    """Each fiscal 2028 month's late-shift level so the golden compliance sits on its target, then the cells rounded
    to one decimal as the board publishes them. Returns {(y, m, dow, h): (forecast, halvorsen)} for those cells and
    the golden compliance the rounded cells give."""
    mu, th = 3600.0 / busy, 3600.0 / pat
    cells, gold = {}, {}
    for y, m in K.fy_months(2028):
        tgt = LV.TARGET[m] / 100.0
        L = LV.solve_level(y, m, LV.TARGET[m], mu, th)
        best = None
        for adj in np.linspace(-0.06, 0.06, 121):
            vals = {}
            for dow in range(7):
                for h in range(24):
                    if is_late_cell(dow, h):
                        vals[(dow, h)] = round((L + adj) * LV.cell_factor(dow, h), 1)
            g = M.month_compliance(y, m, lambda d, hh: vals[(d, hh)], mu, th)[0]
            if best is None or abs(g - tgt) < abs(best[0] - tgt):
                best = (g, vals)
        gold[(y, m)] = best[0]
        for (dow, h), v in best[1].items():
            cells[(y, m, dow, h)] = (v, round(LV.HALV_LATE[m] * LV.cell_factor(dow, h), 1))
    return cells, gold


def tables(busy, pat):
    """{sheet: [row tuples]} and the golden compliance by month."""
    late, gold = centre_late(busy, pat)
    out = {}
    for fy in (2027, 2028):
        rows = []
        for name, scale in PSAPS:
            for y, m in K.fy_months(fy):
                for dow in range(7):
                    for h in range(24):
                        if name == "Tavish County ECA" and fy == 2028 and is_late_cell(dow, h):
                            v, hv = late[(y, m, dow, h)]
                        else:
                            r = K.rng("fc", name, y, m, dow, h)
                            base = base_cell(y, m, dow, h)
                            if name != "Tavish County ECA":
                                base = base * scale * float(r.uniform(0.95, 1.05))
                            hv = 0.0
                            if fy == 2028 and name == "Tavish County ECA":
                                hv = round(halv_cell(m, dow, h), 1)
                            elif fy == 2028 and name == "Wexley County ECC":
                                hv = -round(halv_cell(m, dow, h) * 1.04, 1)
                            v = round(base + hv, 1)
                        rows.append((name, "FY%d" % fy, "%s %d" % (K.MON3[m], y), K.DOW[dow], "%02d:00" % h, v,
                                     hv if hv else 0.0))
        out["FY%d" % fy] = rows
    return out, gold, late


HEAD = ["PSAP", "Fiscal year", "Month", "Day", "Hour beginning", "Forecast 911 calls per hour",
        "Of which Halvorsen move"]


def write(path, busy, pat):
    import openpyxl
    from openpyxl.styles import Alignment, Font
    tb, gold, late = tables(busy, pat)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Read me"
    for i, line in enumerate(READ_ME, start=1):
        ws.cell(row=i, column=1, value=line if line else None)
        if i <= 2:
            ws.cell(row=i, column=1).font = Font(bold=True, size=12 if i == 1 else 11)
        else:
            ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 110
    for sheet in ("FY2027", "FY2028"):
        w = wb.create_sheet(sheet)
        w.cell(row=1, column=1, value="Tri-County Regional 911 Board, hourly 911 volume forecast, %s (July %d to June %d)"
               % (sheet, int(sheet[2:]) - 1, int(sheet[2:]))).font = Font(bold=True)
        for j, h in enumerate(HEAD, start=1):
            w.cell(row=3, column=j, value=h).font = Font(bold=True)
        for i, row in enumerate(tb[sheet], start=4):
            for j, v in enumerate(row, start=1):
                w.cell(row=i, column=j, value=v)
        for col, wd in zip("ABCDEFG", (20, 11, 10, 11, 14, 26, 24)):
            w.column_dimensions[col].width = wd
        w.freeze_panes = "A4"
    wb.save(path)
    return tb, gold, late
