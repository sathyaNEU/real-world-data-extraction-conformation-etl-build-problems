"""Write task167/metadata.json from the built pack: the distractors, the input gates and every shipped file's source,
date and licence."""
from __future__ import annotations

import csv
import json
import os

from openpyxl import load_workbook

import common as C

ORIGIN = "Constructed for this task (seeded, deterministic): "
SOURCES = {
    C.F_REGISTER: ("Brandvold Foundation investment office, fund commitments register, one row per vehicle", "2027-01-11"),
    C.F_ACTIVITY: ("the Foundation's current fund administrator, investor capital activity extract, notice lines with "
                   "value dates 1 April 2016 to 31 December 2026", "2027-01-13"),
    C.F_LEGACY: ("the Foundation's former fund administrator, notice history, payment dates 3 January 2006 to 30 June "
                 "2016", "2016-07-08"),
    C.F_FX: ("the Foundation's custodian, euro reference rates, business days 2006 to 2026 (a constructed series, not an "
             "official release)", "2027-01-04"),
    C.F_POLICY: ("Brandvold Foundation board of directors, liquidity and reserve policy adopted 12 June 2025", "2025-06-16"),
    C.F_CAPACC: ("the Foundation's current fund administrator, capital account summary at 31 December 2026", "2027-01-13"),
    C.F_POS: ("the Foundation's custodian, positions report for accounts 6011-2207 and 6011-2215 at 31 December 2026",
              "2027-01-04"),
    C.F_CUSTACT: ("the Foundation's custodian, settled activity report for account 6011-2207, 2026", "2027-01-04"),
    C.F_DICT: ("Brandvold Foundation investment office, data dictionary", "2027-01-14"),
    C.F_MANIFEST: ("Brandvold Foundation investment office, export manifest for the January 2027 board set", "2027-01-14"),
    C.F_GRANTS: ("Brandvold Foundation grants office, 2027 grants budget memo", "2026-12-11"),
    C.F_PERF: ("general partners' quarterly reports at 30 September 2026, compiled by the investment office", "2026-12-18"),
}
TITLE = ("How large a call reserve a Minnesota grant-making foundation's liquidity policy requires for 2027 on its "
         "private fund commitments")


def _rows(path):
    ext = path.rsplit(".", 1)[-1]
    if ext == "csv":
        with open(path, encoding="utf-8") as fh:
            return sum(1 for _ in csv.reader(fh)) - 1
    if ext == "xlsx":
        wb = load_workbook(path, read_only=True)
        return sum(max(0, ws.max_row - 1) for ws in wb.worksheets)
    return None


def write(task_dir):
    target = os.path.join(task_dir, "target")
    files = []
    big, big_rows = None, 0
    for f in sorted(os.listdir(target)):
        p = os.path.join(target, f)
        rows = _rows(p)
        if rows and rows > big_rows:
            big, big_rows = f, rows
        src, date = SOURCES[f]
        files.append(dict(path=f, format=f.rsplit(".", 1)[-1], bytes=os.path.getsize(p), rows=rows,
                          source=ORIGIN + src, date=date, license="CC0-1.0 (original work, constructed for this task)"))
    meta = dict(
        task="task167", title=TITLE,
        domain="General Statistics & Data Science", client_domain="General Statistics & Data Science",
        house_domain="nonprofit-grant-making",
        subdomain="endowment management (a grant-making foundation's private-markets programme: the year's call reserve "
                  "forecast fund by fund from the call schedules of the funds whose first year was financed like each "
                  "fund's own)",
        objective="Forecasting & Predictive Modeling",
        prompt_shape="one committed reserve with the year's expected calls in total and by strategy, the private "
                     "funds' net asset value at year end by strategy, currency, vintage band and in total, and the "
                     "reserve account's qualifying assets with their bills and the face value maturing each quarter",
        as_of="2027-01-14",
        deliverables=["call_reserve_2027.xlsx", "call_reserve_memo_2027.pdf"],
        distractor_files=list(C.DISTRACTORS),
        input_gates=dict(files=len(files), formats=sorted(set(x["format"] for x in files)), largest_file=big,
                         largest_file_rows=big_rows),
        sources=[dict(files="all files under target/",
                      origin="constructed for this task by task167/generator/build.py (seeded, deterministic) around a "
                             "fictional private foundation in Saint Paul, Minnesota, its fund managers, its two fund "
                             "administrators and its custodian; no real organisation's records or person is depicted.",
                      date="2027-01-14", license="CC0-1.0 (original work)")],
        created="2026-10-18", generator="task167/generator/build.py", files=files)
    with open(os.path.join(task_dir, "metadata.json"), "w") as fh:
        json.dump(meta, fh, indent=1)
        fh.write("\n")
    return meta
