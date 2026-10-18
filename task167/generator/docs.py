"""task167 documents: the board's liquidity and reserve policy (PDF) and the grants office's 2027 budget (Word).
Their text is held here as data so checks.py can quote it and verify.py can read it back."""
from __future__ import annotations

import datetime as dt
import os

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                        # noqa: E402
from reportlab.lib.pagesizes import letter                                              # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                         # noqa: E402
from reportlab.lib.units import mm                                                      # noqa: E402
from reportlab.pdfbase import pdfmetrics                                                # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                            # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import common as C          # noqa: E402
import postprocess as PP    # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/"
for _n, _f in (("LibSerif", "liberation/LiberationSerif-Regular.ttf"), ("LibSerif-Bold", "liberation/LiberationSerif-Bold.ttf"),
               ("Carlito", "crosextra/Carlito-Regular.ttf"), ("Carlito-Bold", "crosextra/Carlito-Bold.ttf")):
    pdfmetrics.registerFont(TTFont(_n, FONT_DIR + _f))
from reportlab.pdfbase.pdfmetrics import registerFontFamily   # noqa: E402
registerFontFamily("LibSerif", normal="LibSerif", bold="LibSerif-Bold", italic="LibSerif", boldItalic="LibSerif-Bold")
registerFontFamily("Carlito", normal="Carlito", bold="Carlito-Bold", italic="Carlito", boldItalic="Carlito-Bold")
INK = colors.HexColor("#1d1f22")
GREY = colors.HexColor("#5b6068")
ACCENT = colors.HexColor("#2b3a55")
RULE = colors.HexColor("#b9bfc8")

POLICY_DATE = dt.date(2025, 6, 16)

POLICY_TITLE = "Liquidity and Reserve Policy"
POLICY = [
    ("1. Purpose", [
        "1.1 This policy sets how the Foundation keeps the liquidity to meet the capital calls of the private funds it "
        "has committed to without selling other assets at a time it does not choose. It applies to the endowment "
        "and to every fund commitment the Foundation holds directly or through a parallel vehicle.",
    ]),
    ("2. The call reserve", [
        "2.1 On the first business day of each year the Foundation holds in its reserve account qualifying assets at "
        "least equal to the call reserve for that year.",
        "2.2 The call reserve for a year is the capital calls the Foundation expects to pay in that calendar year on "
        "the fund commitments it made before the year began, rounded up to the next $50 million.",
        "2.3 Capital calls include the management fees, partnership expenses and organisational expenses a fund calls "
        "with them, and are measured before any distribution a fund sets against a call. Commitments the Foundation "
        "makes during the year are provided for from the endowment when they are made and are not part of the call "
        "reserve.",
        "2.4 For the call reserve, a fund's age is counted in months from its first close.",
        "2.5 Amounts in euros are converted at the custodian's reference rate for the last business day of the prior "
        "year.",
    ]),
    ("3. Qualifying assets", [
        "3.1 Qualifying assets are, at market value on a settled basis: United States Treasury bills maturing on or "
        "before 31 December of the reserve year; shares of the government money market fund named in Schedule 1; and "
        "cash in the reserve account.",
        "3.2 Other holdings, including any held in the Foundation's operating account, do not count toward the call "
        "reserve.",
    ]),
    ("4. Valuing the private funds", [
        "4.1 For the board's liquidity report each fund is carried at its most recent reported net asset value, "
        "adjusted for the contributions paid and distributions received after the date of that value, and euro values "
        "are converted at the rate in 2.5.",
    ]),
    ("5. Certification and review", [
        "5.1 The director of investment risk certifies the call reserve and the qualifying assets to the board at its "
        "January meeting. Any shortfall is met by a transfer from the endowment pool within ten business days of the "
        "meeting.",
        "5.2 The board reviews this policy every three years.",
    ]),
]
def section_numbers():
    """the policy's numbered sections, as the write-up cites them"""
    return [p.split(" ", 1)[0] for _, paras in POLICY for p in paras if p[:1].isdigit()]


SCHEDULE_1 = "Schedule 1. Approved money market fund: Ironvale Government Money Market Fund (all share classes)."
ADOPTED = "Adopted by the board of directors on 12 June 2025. Replaces the liquidity guidelines of March 2019."


def _styles():
    return dict(
        h1=ParagraphStyle("h1", fontName="LibSerif-Bold", fontSize=15.5, leading=19, textColor=ACCENT, spaceAfter=3),
        sub=ParagraphStyle("sub", fontName="Carlito", fontSize=9.2, leading=12, textColor=GREY, spaceAfter=10),
        h2=ParagraphStyle("h2", fontName="LibSerif-Bold", fontSize=11.2, leading=14, textColor=ACCENT, spaceBefore=7,
                          spaceAfter=3),
        body=ParagraphStyle("body", fontName="LibSerif", fontSize=10.4, leading=14.2, textColor=INK, spaceAfter=5),
        small=ParagraphStyle("small", fontName="Carlito", fontSize=8.6, leading=11, textColor=GREY, spaceAfter=3))


def _footer(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Carlito", 7.8)
        canvas.setFillColor(GREY)
        canvas.drawString(24 * mm, 11 * mm, text)
        canvas.drawRightString(letter[0] - 24 * mm, 11 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def pdf_stamp(d, hh=10, mi=0):
    off = C.us_central_offset(dt.datetime.combine(d, dt.time(hh, mi)))
    return ("D:%04d%02d%02d%02d%02d00%+03d'00'" % (d.year, d.month, d.day, hh, mi, off)).encode()


def policy_text():
    out = [POLICY_TITLE, C.ORG]
    for h, paras in POLICY:
        out.append(h)
        out.extend(paras)
    out.extend([SCHEDULE_1, ADOPTED])
    return out


def write_policy(out):
    path = os.path.join(out, C.F_POLICY)
    st = _styles()
    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=24 * mm, rightMargin=24 * mm, topMargin=22 * mm,
                            bottomMargin=20 * mm, title=POLICY_TITLE, author=C.ORG, subject="", creator="", producer="")
    story = [Paragraph(esc(C.ORG), st["sub"]), Paragraph(esc(POLICY_TITLE), st["h1"]),
             Paragraph(esc(ADOPTED), st["sub"])]
    for h, paras in POLICY:
        story.append(Paragraph(esc(h), st["h2"]))
        for p in paras:
            story.append(Paragraph(esc(p), st["body"]))
    story.append(Spacer(1, 6))
    story.append(Paragraph(esc(SCHEDULE_1), st["body"]))
    doc.build(story, onFirstPage=_footer("Brandvold Foundation, board policy B-7"),
              onLaterPages=_footer("Brandvold Foundation, board policy B-7"))
    PP.normalize_pdf(path, pdf_stamp(POLICY_DATE, 10, 12), pdf_stamp(POLICY_DATE, 10, 12))


GRANTS_PARAS = [
    "To: Board of directors",
    "From: Grants office",
    "Date: 11 December 2026",
    "Re: 2027 grants budget as approved on 10 December 2026",
    "The board approved grants of $206.4 million for 2027, against $198.1 million budgeted for 2026. The budget meets "
    "the Foundation's distribution target of 5.2 per cent of the average endowment value over the twelve quarters to "
    "30 September 2026, and the federal minimum distribution with room to spare.",
    "By programme the budget is: education and early learning $61.8 million; health and community wellbeing $48.3 "
    "million; arts and culture $29.7 million; environment and the Mississippi watershed $33.9 million; housing and "
    "economic opportunity $26.1 million; and discretionary and matching grants $6.6 million.",
    "Payments follow the grant agreements' schedules. We expect to pay $47.2 million in the first quarter, $58.9 "
    "million in the second, $44.6 million in the third and $55.7 million in the fourth, from the operating account, "
    "which the endowment pool funds monthly.",
    "Multi-year grants already awarded account for $121.5 million of the 2027 total; the remaining $84.9 million will be "
    "awarded during the year under the programme guidelines the board adopted in September.",
]


def write_grants(out):
    from docx import Document
    from docx.shared import Pt
    path = os.path.join(out, C.F_GRANTS)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    doc.add_heading("2027 grants budget", level=1)
    for p in GRANTS_PARAS:
        doc.add_paragraph(p)
    doc.save(path)
    created, modified = C.STAMPS[C.F_GRANTS]
    PP.normalize_docx(path, "Adriana Mitchell", "Adriana Mitchell", created, modified, revision=6, total_time=94, seed=167,
                      title="2027 grants budget", company=C.ORG)
