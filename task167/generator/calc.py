"""task167 calc: every graded figure, rung and rival computed from the shipped files alone (the target folder),
never from generator state. checks.py asserts on these numbers; verify.py recomputes them on its own code path.

    python3 calc.py TARGET_DIR"""
from __future__ import annotations

import csv
import datetime as dt
import math
import os
import sys
from collections import defaultdict

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import common as C   # noqa: E402

CALL_LINES = ("Investments", "Management fees", "Management fee offset", "Partnership expenses",
              "Organisational expenses")
DIST_LINES = ("Return of capital", "Realised gain", "Income")
YEAR = 2027
HANDOVER = dt.date(2016, 4, 1)
YE = dt.date(2026, 12, 31)
MMAX = 252


def mi(d):
    return (d.year - 2006) * 12 + d.month - 1


def _d(x):
    if isinstance(x, dt.datetime):
        return x.date()
    if isinstance(x, dt.date):
        return x
    return dt.date.fromisoformat(str(x)[:10])


def norm(s):
    return " ".join(str(s).upper().replace(",", "").replace(".", "").split())


class Pack:
    """the shipped files, read once"""

    def __init__(self, target):
        self.t = target
        self.read_register()
        self.read_fx()
        self.read_activity()
        self.read_legacy()
        self.build_calls()

    def read_register(self):
        wb = load_workbook(os.path.join(self.t, C.F_REGISTER), read_only=True, data_only=True)
        ws = wb["Commitments"]
        rows = list(ws.iter_rows(values_only=True))
        h = rows[0]
        self.veh = {}
        for r in rows[1:]:
            d = dict(zip(h, r))
            if not d["Fund ID"]:
                continue
            d["First close"] = _d(d["First close"])
            d["fund"] = d["Fund ID"].rstrip("P")
            self.veh[d["Administrator code"]] = d
        self.by_legacy = {norm(d["Former administrator name"]): code for code, d in self.veh.items()
                          if d["Former administrator name"]}

    def read_fx(self):
        self.fx = {}
        with open(os.path.join(self.t, C.F_FX), encoding="utf-8") as fh:
            for r in csv.DictReader(fh):
                self.fx[dt.date.fromisoformat(r["date"])] = float(r["usd_per_eur"])
        self.ye_rate = self.fx[max(d for d in self.fx if d <= YE)]

    def read_activity(self):
        self.lines = []
        with open(os.path.join(self.t, C.F_ACTIVITY), encoding="utf-8") as fh:
            for r in csv.DictReader(fh):
                if r["status"] != "Settled":
                    continue
                self.lines.append((r["fund_code"], dt.date.fromisoformat(r["value_date"]), r["line_type"],
                                   float(r["amount"]), r["notice_ref"], r["notice_type"]))

    def read_legacy(self):
        wb = load_workbook(os.path.join(self.t, C.F_LEGACY), read_only=True, data_only=True)
        rows = list(wb["Notices"].iter_rows(values_only=True))
        h = rows[0]
        self.legacy = []
        for r in rows[1:]:
            d = dict(zip(h, r))
            if not d["Notice No."]:
                continue
            d["Payment Date"] = _d(d["Payment Date"])
            d["code"] = self.by_legacy[norm(d["Fund"])]
            self.legacy.append(d)

    def build_calls(self):
        """calls by vehicle and month: the legacy file before the handover, the current file from 1 April 2016"""
        self.calls = defaultdict(lambda: defaultdict(float))
        self.dists = defaultdict(lambda: defaultdict(float))
        for code, vd, lt, amt, ref, nt in self.lines:
            if lt in CALL_LINES:
                self.calls[code][mi(vd)] += amt
            elif lt in DIST_LINES:
                self.dists[code][mi(vd)] += -amt
        for d in self.legacy:
            if d["Payment Date"] >= HANDOVER:
                continue
            c = d["Investments"] + d["Management Fees"] + d["Fund Expenses"] + d["Organisation Costs"]
            x = d["Return of Capital"] + d["Capital Gains"] + d["Income"]
            self.calls[d["code"]][mi(d["Payment Date"])] += c
            self.dists[d["code"]][mi(d["Payment Date"])] += x


class Book:
    """fund-level view (main and parallel vehicles together) for the call forecast"""

    def __init__(self, p, clock="month"):
        self.p = p
        funds = defaultdict(list)
        for code, d in p.veh.items():
            funds[d["fund"]].append(code)
        self.funds = {}
        asof_m = mi(dt.date(2027, 1, 1))
        for fid, codes in funds.items():
            d0 = p.veh[codes[0]]
            fc = d0["First close"]
            m0 = mi(fc)
            com = sum(p.veh[c]["Commitment"] for c in codes)
            h = [0.0] * MMAX
            for c in codes:
                for m, a in p.calls[c].items():
                    if m < asof_m:
                        h[m - m0] += a
            obs = asof_m - m0
            self.funds[fid] = dict(fid=fid, codes=codes, strategy=d0["Strategy"], vintage=d0["Vintage"], ccy=d0["Currency"],
                                   commitment=com, first_close=fc, m0=m0, h=h, obs=obs, paid=sum(h),
                                   programme=fc >= dt.date(2026, 1, 1), status=d0["Status"])
            n1 = min(12, obs)
            self.funds[fid]["y1"] = sum(h[:n1]) / com * (12.0 / n1) if n1 > 0 else None

    def fx_of(self, f):
        return self.p.ye_rate if f["ccy"] == "EUR" else 1.0


def curve(book, sel, step=1, rate=False, m_from=None, weight="commitment"):
    """the call schedule by age bin of `step` months: commitment-weighted share of commitment (or, with rate, the share
    of the undrawn balance), over the funds that have been observed through each bin"""
    nb = MMAX // step
    num, den = [0.0] * nb, [0.0] * nb
    for f in book.funds.values():
        if not sel(f):
            continue
        cum = 0.0
        for b in range(nb):
            if (b + 1) * step > f["obs"]:
                break
            x = sum(f["h"][b * step:(b + 1) * step])
            if m_from is None or f["m0"] + b * step >= m_from:
                if weight == "equal":
                    num[b] += x / f["commitment"]
                    den[b] += 1.0 if not rate else (f["commitment"] - cum) / f["commitment"]
                else:
                    num[b] += x
                    den[b] += (f["commitment"] - cum) if rate else f["commitment"]
            cum += x
    return [n / d if d > 0 else 0.0 for n, d in zip(num, den)]


def forecast(book, curvef, step=1, rate=False, by=None):
    """expected calls in calendar 2027 by fund, in US dollars at the year-end reference rate"""
    out = {}
    for fid, f in book.funds.items():
        if f["status"] != "Active":
            out[fid] = 0.0
            continue
        cv = curvef(f)
        rem = f["commitment"] - f["paid"]
        got = 0.0
        for k in range(12):
            a = f["obs"] + k
            b = a // step
            if a < 0:
                continue
            if b >= len(cv):
                break
            if rate:
                c = (rem - got) * (1.0 - (1.0 - min(cv[b], 0.999999)) ** (1.0 / step))
            else:
                c = f["commitment"] * cv[b] / step
            c = max(0.0, min(c, rem - got))
            got += c
        out[fid] = got * book.fx_of(f)
    return out


def is_line(f, t=0.06):
    return f["y1"] is not None and f["y1"] < t


def hist(f):
    return not f["programme"]


def golden_call(book, step=1, rate=False, thr=0.06, line_from=None, weight="commitment"):
    sel_b = (lambda f: hist(f) and is_line(f, thr) and (line_from is None or f["vintage"] >= line_from))
    sel_a = (lambda f: hist(f) and not is_line(f, thr))
    cb = curve(book, sel_b, step, rate, weight=weight)
    ca = curve(book, sel_a, step, rate, weight=weight)
    return forecast(book, lambda f: cb if is_line(f, thr) else ca, step, rate)


RESERVE_STEP_MUSD = 50          # policy 2.2: rounded up to the next $50 million


def reserve(total_usd):
    return int(math.ceil(total_usd / (RESERVE_STEP_MUSD * 1e6) - 1e-9) * RESERVE_STEP_MUSD)


def summarise(book, fc):
    tot = sum(fc.values())
    by_s = defaultdict(float)
    prog = 0.0
    for fid, v in fc.items():
        by_s[book.funds[fid]["strategy"]] += v
        if book.funds[fid]["programme"]:
            prog += v
    return dict(total=tot, by_strategy=dict(by_s), programme=prog, earlier=tot - prog, reserve=reserve(tot))


def main_call(target):
    p = Pack(target)
    book = Book(p)
    g = summarise(book, golden_call(book))
    return p, book, g


# ------------------------------------------------------------------------------------------- Ask A, NAV at year end
def read_capital_accounts(target):
    wb = load_workbook(os.path.join(target, C.F_CAPACC), read_only=True, data_only=True)
    rows = list(wb["Summary"].iter_rows(values_only=True))
    hi = [i for i, r in enumerate(rows) if r and r[0] == "Fund code"][0]
    h = rows[hi]
    out = []
    for r in rows[hi + 1:]:
        if not r or not r[0]:
            continue
        d = dict(zip(h, r))
        d["NAV date"] = _d(d["NAV date"]) if d["NAV date"] else None
        out.append(d)
    return out


def nav_year_end(p, target, combined_seen=True, currency_from_register=True, roll=True, fx_at_nav_date=False):
    """each fund at its latest reported NAV rolled forward for contributions and distributions to 31 December 2026,
    euro values at the year-end reference rate. A main vehicle's row is the general partner's combined statement for
    the main and parallel interests (combined_seen=False sums every row, the parallel's counted twice); the fund's
    currency is the register's (currency_from_register=False takes the summary's label)."""
    ca = read_capital_accounts(target)
    out = []
    for d in ca:
        code = d["Fund code"]
        v = p.veh[code]
        nav0, nd = (0.0, None) if d["Reported NAV"] is None else (d["Reported NAV"], d["NAV date"])
        if v["Vehicle"] == "Parallel" and combined_seen:
            continue
        codes = [code]
        if v["Vehicle"] == "Main" and combined_seen:
            codes = [c for c, x in p.veh.items() if x["fund"] == v["fund"]]
        adj = 0.0
        if roll:
            start = mi(nd) + 1 if nd else -1
            for c in codes:
                adj += sum(a for m, a in p.calls[c].items() if start <= m <= mi(YE))
                adj -= sum(a for m, a in p.dists[c].items() if start <= m <= mi(YE))
        val = nav0 + adj
        ccy = v["Currency"] if currency_from_register else d["Currency"]
        rate = 1.0
        if ccy == "EUR":
            rate = p.ye_rate if not fx_at_nav_date else p.fx[max(x for x in p.fx if x <= (nd or YE))]
        out.append(dict(code=code, fund=v["fund"], strategy=v["Strategy"], vintage=v["Vintage"], ccy=ccy, usd=val * rate))
    return out


def nav_figures(rows):
    f = defaultdict(float)
    for r in rows:
        f["total"] += r["usd"]
        f["strategy:" + r["strategy"]] += r["usd"]
        f["ccy:" + r["ccy"]] += r["usd"]
    return dict(f)


# ------------------------------------------------------------------------------ Ask B, qualifying reserve assets
def read_positions(target):
    with open(os.path.join(target, C.F_POS), encoding="utf-8") as fh:
        return [r for r in csv.DictReader(fh) if r["asset_id"]]


def read_custody_activity(target):
    with open(os.path.join(target, C.F_CUSTACT), encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def settled_holdings(target):
    """the reserve account's holdings at 31 December 2026 rebuilt from its opening position and settled activity"""
    q = defaultdict(float)
    for r in read_custody_activity(target):
        if r["account"] != "6011-2207" or r["asset_id"] == "CASH":
            continue
        qty = float(r["quantity"]) if r["quantity"] else 0.0
        if r["transaction"] in ("OPENING POSITION", "PURCHASE", "DIVIDEND REINVEST"):
            q[r["asset_id"]] += qty
        elif r["transaction"] in ("MATURITY", "REDEMPTION"):
            q[r["asset_id"]] += qty      # negative quantities
    return {k: round(v, 2) for k, v in q.items() if abs(v) > 0.005}


def qualifying(target, settled=True, by_description=True):
    """the reserve account's qualifying assets at 31 December 2026 under policy 3.1: Treasury bills maturing in 2027,
    the Schedule 1 government money market fund and the account's cash, at market value on a settled basis.
    settled=False takes the trade-date positions report as it stands; by_description=False takes the custodian's
    "Government Bill" class as Treasury bills."""
    pos = read_positions(target)
    held = settled_holdings(target)
    out = dict(bills=0.0, mmf=0.0, cash=0.0, ladder=[0.0, 0.0, 0.0, 0.0])
    for r in pos:
        if r["account"] != "6011-2207":
            continue
        kind, mv, qty = r["asset_type"], float(r["market_value_usd"]), float(r["quantity"])
        if kind == "Cash":
            out["cash"] += mv
            continue
        if settled and r["asset_id"] not in held:
            continue
        if kind == "Government Bill":
            if by_description and not r["description"].startswith("U.S. TREASURY BILL"):
                continue
            mat = dt.date.fromisoformat(r["maturity_date"])
            if mat > dt.date(2027, 12, 31):
                continue
            out["bills"] += mv
            out["ladder"][(mat.month - 1) // 3] += qty
        elif kind == "Money Market Fund":
            if "Government Money Market Fund" in r["description"]:
                out["mmf"] += mv
    out["total"] = out["bills"] + out["mmf"] + out["cash"]
    return out


if __name__ == "__main__":
    t = sys.argv[1]
    p, book, g = main_call(t)
    print("reserve", g["reserve"], "total %.1f programme %.1f earlier %.1f" % (g["total"] / 1e6, g["programme"] / 1e6,
                                                                               g["earlier"] / 1e6))
    print({k: round(v / 1e6, 1) for k, v in g["by_strategy"].items()})
    nf = nav_figures(nav_year_end(p, t))
    print({k: round(v / 1e6, 1) for k, v in sorted(nf.items())})
    q = qualifying(t)
    print({k: (round(v / 1e6, 2) if not isinstance(v, list) else [round(x / 1e6, 2) for x in v]) for k, v in q.items()})
    sys.stdout.flush()
