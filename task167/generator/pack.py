"""task167 pack: writes every file of the evidence pack into a folder.

    python3 pack.py OUT_DIR"""
from __future__ import annotations

import datetime as dt
import os
import sys

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import common as C          # noqa: E402
import custody as CU        # noqa: E402
import docs as D            # noqa: E402
import notices as N         # noqa: E402
import postprocess as PP    # noqa: E402
import records as R         # noqa: E402
import world as W           # noqa: E402

LEGACY_TYPE = {"call": "Capital Call", "dist": "Distribution", "net": "Call/Distribution (net)"}
CUR_TYPE = {"call": "Capital call", "dist": "Distribution", "net": "Capital call and distribution"}
LINE_ORDER = ("invest", "fee", "offset", "expense", "org", "roc", "gain", "income")
NAV_DATE = {"Buyout": dt.date(2026, 9, 30), "Growth equity": dt.date(2026, 9, 30),
            "Venture capital": dt.date(2026, 6, 30), "Private credit": dt.date(2026, 11, 30)}


def legacy_name(name):
    s = name.upper().replace(",", "").replace(".", "")
    s = s.replace("CAPITAL PARTNERS", "CAP PTRS").replace("PARTNERS", "PTRS").replace("GROWTH FUND", "GROWTH FD")
    s = s.replace(" FUND ", " FD ").replace("DIRECT LENDING", "DIR LENDING").replace("VENTURES", "VENT")
    return s.replace(" LP", " LP").strip()


class World:
    """everything the writers need, built once"""

    def __init__(self):
        self.funds, self.recs, self.fx = R.all_records()
        self.vehicles = [v for f in self.funds for v in f.vehicles]
        self.notices = N.build_notices(self.funds, self.recs)
        self.cancelled = N.cancellations(self.notices)
        N.assign_refs(self.notices, self.cancelled)
        self.fund_id = {}
        k = 101
        for f in sorted(self.funds, key=lambda f: (f.vintage, f.first_close, f.name)):
            self.fund_id[f.id] = "BF-%04d" % k
            k += 1 + (f.id * 7) % 3
        self.status = {}
        for f in self.funds:
            nav_ye = self.recs[f.id][1].get(W.month_index(W.AS_OF), 0.0)
            open_deals = any(d[0] <= W.month_index(W.AS_OF) < d[2] for d in self.recs[f.id][2])
            self.status[f.id] = "Active" if (open_deals or nav_ye > 0.0005 or f.vintage >= 2016) else "Realised"
        self.custody = CU.build(self.notices, self.fx)


# ------------------------------------------------------------------------------------------------ office helpers
HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")


def _sheet(ws, head, rows, widths=None, start=1, num_fmt=None):
    for j, h in enumerate(head):
        c = ws.cell(row=start, column=j + 1, value=h)
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = ws.cell(row=start + 1 + i, column=j + 1, value=v)
            if num_fmt and j in num_fmt and isinstance(v, (int, float)):
                c.number_format = num_fmt[j]
            if isinstance(v, dt.date):
                c.number_format = "yyyy-mm-dd"
    if widths:
        for j, w in enumerate(widths):
            ws.column_dimensions[chr(65 + j) if j < 26 else "A" + chr(65 + j - 26)].width = w
    ws.freeze_panes = ws.cell(row=start + 1, column=1)


def _save(wb, path, creator, last_by, key, company=C.ORG):
    wb.save(path)
    created, modified = C.STAMPS[key]
    PP.normalize_ooxml(path, creator, last_by, created, modified, excel=True, company=company)


# ------------------------------------------------------------------------------------------------------ the files
def write_register(out, w):
    wb = Workbook()
    ws = wb.active
    ws.title = "Commitments"
    rows = []
    for f in sorted(w.funds, key=lambda f: (f.vintage, f.first_close, f.name)):
        fid = w.fund_id[f.id]
        for v in f.vehicles:
            legacy = legacy_name(v.name) if f.first_close <= W.LEGACY_CUTOFF else None
            note = None
            if v.parallel:
                note = "Parallel vehicle of %s; called and distributed pro rata with the main vehicle" % fid
            elif len(f.vehicles) > 1:
                note = "Main vehicle; parallel vehicle %sP" % fid
            rows.append((fid + ("P" if v.parallel else ""), v.name, "Parallel" if v.parallel else "Main", v.code, legacy,
                         f.manager, f.strategy, f.vintage, f.first_close, f.ccy, round(v.commitment * 1e6, 2),
                         w.status[f.id], note))
    _sheet(ws, ("Fund ID", "Fund", "Vehicle", "Administrator code", "Former administrator name", "Manager", "Strategy",
                "Vintage", "First close", "Currency", "Commitment", "Status", "Notes"), rows,
           widths=(10, 44, 9, 15, 34, 30, 15, 8, 11, 9, 16, 9, 60), num_fmt={10: "#,##0.00"})
    _save(wb, os.path.join(out, C.F_REGISTER), "Kristin Brown", "Brian Erickson", C.F_REGISTER)
    return len(rows)


def activity_rows(w):
    rows = []
    allc = [n for n in w.notices if n.value_date >= W.HANDOVER_START] + list(w.cancelled)
    allc.sort(key=lambda n: (n.value_date, n.notice_date, n.ref))
    k = 0
    for n in allc:
        for kind in LINE_ORDER:
            if kind not in n.lines:
                continue
            amt = n.lines[kind]
            sign = 1 if kind in N.CALL_KINDS else -1
            k += 1
            rows.append(("A%07d" % (1340000 + k * 7 + (k % 5)), n.ref, n.vehicle.code, CUR_TYPE[n.kind],
                         C.d_iso(n.notice_date), C.d_iso(n.value_date), N.LINE_NAME[kind], C.money(sign * amt),
                         n.vehicle.fund.ccy, n.status))
    return rows


def write_activity(out, w):
    rows = activity_rows(w)
    return C.write_csv(os.path.join(out, C.F_ACTIVITY),
                       ("activity_id", "notice_ref", "fund_code", "notice_type", "notice_date", "value_date", "line_type",
                        "amount", "currency", "status"), rows)


def write_legacy(out, w):
    wb = Workbook()
    ws = wb.active
    ws.title = "Notices"
    rows = []
    for n in sorted([n for n in w.notices if n.value_date <= W.LEGACY_CUTOFF], key=lambda n: (n.notice_date, n.legacy_ref)):
        g = lambda k: round(n.lines.get(k, 0.0), 2)
        calls = sum(g(k) for k in N.CALL_KINDS)
        dists = sum(g(k) for k in N.DIST_KINDS)
        rows.append((n.legacy_ref, legacy_name(n.vehicle.name), n.notice_date, n.value_date, LEGACY_TYPE[n.kind],
                     g("invest"), round(g("fee") + g("offset"), 2), g("expense"), g("org"), g("roc"), g("gain"), g("income"),
                     round(calls - dists, 2), n.vehicle.fund.ccy))
    _sheet(ws, ("Notice No.", "Fund", "Notice Date", "Payment Date", "Type", "Investments", "Management Fees",
                "Fund Expenses", "Organisation Costs", "Return of Capital", "Capital Gains", "Income", "Net Amount",
                "Ccy"), rows, widths=(10, 40, 12, 12, 22, 15, 15, 13, 13, 15, 15, 12, 15, 6),
           num_fmt={i: "#,##0.00" for i in range(5, 13)})
    _save(wb, os.path.join(out, C.F_LEGACY), "Fund Services", "Fund Services", C.F_LEGACY, company="")
    return len(rows)


def write_fx(out, w):
    rows = [(C.d_iso(d), "%.4f" % r) for d, r in sorted(w.fx.items())]
    return C.write_csv(os.path.join(out, C.F_FX), ("date", "usd_per_eur"), rows)


def capital_accounts(w):
    """one row per vehicle with a position at 31 December 2026; a main vehicle's row carries the general partner's
    combined statement for the Foundation's main and parallel interests"""
    rows = []
    hist_calls, hist_dists = {}, {}
    for n in w.notices:
        if n.status != "Settled":
            continue
        hist_calls[n.vehicle.code] = hist_calls.get(n.vehicle.code, 0.0) + n.calls()
        hist_dists[n.vehicle.code] = hist_dists.get(n.vehicle.code, 0.0) + n.dists()
    for f in sorted(w.funds, key=lambda f: (f.strategy, f.vintage, f.name)):
        if w.status[f.id] != "Active":
            continue
        nd = NAV_DATE[f.strategy]
        nav_f = w.recs[f.id][1].get(W.month_index(nd), 0.0) * 1e6
        for v in f.vehicles:
            if v.parallel:
                vs = [v]
            else:
                vs = f.vehicles          # the combined statement
            com = sum(x.commitment for x in vs) * 1e6
            con = sum(hist_calls.get(x.code, 0.0) for x in vs)
            dis = sum(hist_dists.get(x.code, 0.0) for x in vs)
            nav = round(nav_f * sum(x.share for x in vs), 2)
            if f.first_close > nd:
                nav, ndv = None, None
            else:
                ndv = nd
            label = "USD" if f.first_close >= dt.date(2021, 1, 1) else f.ccy
            rows.append((v.code, v.name, label, round(com, 2), round(con, 2), round(dis, 2), round(com - con, 2), nav,
                         ndv))
    return rows


def write_capital_accounts(out, w):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.cell(row=1, column=1, value="Capital account summary, Brandvold Foundation").font = Font(bold=True, size=12)
    ws.cell(row=2, column=1, value="Prepared by the fund administrator at 31 December 2026 from the general partners' "
                                  "statements received by 12 January 2027. Amounts in each fund's currency.")
    rows = capital_accounts(w)
    _sheet(ws, ("Fund code", "Fund", "Currency", "Commitment", "Contributions to date", "Distributions to date",
                "Unfunded commitment", "Reported NAV", "NAV date"), rows, widths=(15, 44, 9, 16, 18, 18, 18, 16, 11),
           start=4, num_fmt={i: "#,##0.00" for i in range(3, 8)})
    _save(wb, os.path.join(out, C.F_CAPACC), "Fund Services", "Fund Services", C.F_CAPACC, company="")
    return len(rows)


def write_custody(out, w):
    acts, pos, truth = w.custody
    rows = []
    for (settle, trade, acct, txn, sid, desc, qty, amount, ref) in sorted(acts, key=lambda a: (a[0], a[3] != "OPENING CASH BALANCE", a[3] != "OPENING POSITION")):
        rows.append((C.d_iso(settle), C.d_iso(trade), acct, txn, sid, desc, ("%.2f" % qty) if qty else "",
                     C.money(amount) if amount else "0.00", ref))
    n1 = C.write_csv(os.path.join(out, C.F_CUSTACT),
                     ("settle_date", "trade_date", "account", "transaction", "asset_id", "description", "quantity",
                      "amount_usd", "reference"), rows)
    prow = []
    for acct in (CU.RESERVE_ACCT, CU.OPERATING_ACCT):
        sub = [p for p in pos if p[0] == acct]
        sub.sort(key=lambda p: ({"Government Bill": 0, "Government Note": 1, "Money Market Fund": 3, "Cash": 4}[p[3]],
                                p[7] or dt.date(2100, 1, 1), p[2]))
        for (a, sid, desc, kind, qty, price, mv, mat) in sub:
            prow.append(("2026-12-31", a, CU.ACCT_NAME[a], sid, desc, kind, "%.2f" % qty, "%.6f" % price, C.money(mv),
                         C.d_iso(mat) if mat else ""))
        prow.append(("2026-12-31", acct, CU.ACCT_NAME[acct], "", "ACCOUNT TOTAL", "", "", "",
                     C.money(sum(p[6] for p in sub)), ""))
    n2 = C.write_csv(os.path.join(out, C.F_POS),
                     ("as_of", "account", "account_name", "asset_id", "description", "asset_type", "quantity", "price",
                      "market_value_usd", "maturity_date"), prow)
    return n1, n2


def write_perf(out, w):
    rng = np.random.default_rng([W.SEED, 31])
    rows = []
    for f in sorted(w.funds, key=lambda f: (f.vintage, f.name)):
        if w.status[f.id] != "Active" or f.programme or f.first_close > dt.date(2025, 9, 30):
            continue
        m = W.month_index(dt.date(2026, 9, 30))
        paid = sum(x.amount for x in w.recs[f.id][0] if x.kind in N.CALL_KINDS and x.month <= m)
        dist = sum(x.amount for x in w.recs[f.id][0] if x.kind in N.DIST_KINDS and x.month <= m)
        nav = w.recs[f.id][1].get(m, 0.0)
        if paid <= 0:
            continue
        dpi, rvpi = dist / paid, nav / paid
        tvpi = dpi + rvpi
        yrs = max(0.75, (dt.date(2026, 9, 30) - f.first_close).days / 365.25)
        irr = (max(tvpi, 0.05) ** (1.0 / max(yrs * 0.62, 0.5)) - 1.0) * 100
        q = ("top quartile", "second quartile", "third quartile", "bottom quartile")[min(3, max(0, int(3.2 - tvpi * 1.4)))]
        rows.append((f.vehicles[0].code, f.name, "2026-09-30", "%.1f" % irr, "%.2f" % dpi, q,
                     C.d_iso(dt.date(2026, 11, 2) + dt.timedelta(days=int(rng.integers(0, 40))))))
    return C.write_csv(os.path.join(out, C.F_PERF), ("fund_code", "fund", "quarter_end", "net_irr_pct", "dpi",
                                                     "gp_benchmark_quartile", "received"), rows)


DICT = [
    (C.F_REGISTER, "Fund ID", "The investment office's identifier for a fund; a parallel vehicle carries its main fund's identifier with P."),
    (C.F_REGISTER, "Vehicle", "Main or Parallel. A parallel vehicle is a separate partnership through which part of the Foundation's commitment to a fund is held."),
    (C.F_REGISTER, "Administrator code", "The vehicle's code in the current fund administrator's records."),
    (C.F_REGISTER, "Former administrator name", "The vehicle's name in the former administrator's notice records, for vehicles that held a first close before 1 July 2016."),
    (C.F_REGISTER, "Vintage", "The year of the fund's first close."),
    (C.F_REGISTER, "First close", "The date of the fund's first close, as stated in its limited partnership agreement."),
    (C.F_REGISTER, "Commitment", "The vehicle's commitment in its currency, as in the subscription agreement."),
    (C.F_REGISTER, "Status", "Active, or Realised once the general partner has made its final distribution."),
    (C.F_ACTIVITY, "activity_id", "Unique identifier of one line of one notice."),
    (C.F_ACTIVITY, "notice_ref", "The administrator's reference for the notice. A notice has one line for each type of amount on it."),
    (C.F_ACTIVITY, "fund_code", "The vehicle's administrator code (see the register)."),
    (C.F_ACTIVITY, "notice_type", "Capital call, Distribution, or Capital call and distribution where the general partner set a distribution against a call on one notice."),
    (C.F_ACTIVITY, "notice_date", "The date of the general partner's notice."),
    (C.F_ACTIVITY, "value_date", "The date the amount was paid or received."),
    (C.F_ACTIVITY, "line_type", "Investments, Management fees, Partnership expenses and Organisational expenses are called from the Foundation; Return of capital, Realised gain and Income are distributed to it. A Management fee offset reduces the management fees called on the same notice."),
    (C.F_ACTIVITY, "amount", "In the vehicle's currency. Positive amounts are paid by the Foundation, negative amounts are paid to it."),
    (C.F_ACTIVITY, "status", "Settled, or Cancelled for a notice the general partner withdrew."),
    (C.F_LEGACY, "Notice No.", "The former administrator's notice number."),
    (C.F_LEGACY, "Fund", "The vehicle's name in the former administrator's records."),
    (C.F_LEGACY, "Payment Date", "The date the amount was paid or received."),
    (C.F_LEGACY, "Investments to Income", "Amounts on the notice by type, in the vehicle's currency, each as a positive number."),
    (C.F_LEGACY, "Net Amount", "Amounts called less amounts distributed on the notice."),
    (C.F_FX, "usd_per_eur", "The custodian's reference rate, US dollars per euro, for each business day."),
    (C.F_CAPACC, "Contributions to date", "All amounts called since the fund's first close."),
    (C.F_CAPACC, "Reported NAV", "The Foundation's share of the fund's net asset value in the general partner's most recent statement."),
    (C.F_CAPACC, "NAV date", "The date of the statement the reported NAV is taken from."),
    (C.F_POS, "as_of", "The close the report is struck at. Positions are reported on trade date and include purchases not yet settled; the cash line is the account's settled balance."),
    (C.F_POS, "quantity", "Face value for bills, notes and discount notes; shares for money market funds."),
    (C.F_POS, "market_value_usd", "The custodian's valuation at the close of 31 December 2026."),
    (C.F_CUSTACT, "settle_date", "The date the transaction settled in the account."),
    (C.F_CUSTACT, "amount_usd", "Cash into the account is positive, cash out negative."),
    (C.F_PERF, "net_irr_pct", "Net internal rate of return since inception as reported by the general partner."),
    (C.F_PERF, "dpi", "Distributions to date over contributions to date, as reported by the general partner."),
    (C.F_PERF, "gp_benchmark_quartile", "The quartile the general partner reports against its own benchmark for the vintage."),
]


def write_dictionary(out, w):
    wb = Workbook()
    ws = wb.active
    ws.title = "Fields"
    _sheet(ws, ("File", "Field", "Definition"), DICT, widths=(46, 26, 110))
    _save(wb, os.path.join(out, C.F_DICT), "Brian Erickson", "Troy Stein", C.F_DICT)
    return len(DICT)


MANIFEST = [
    (C.F_REGISTER, "Investment office commitments register", "2027-01-11", "every fund commitment since 2006, one row per vehicle"),
    (C.F_ACTIVITY, "Current fund administrator, investor activity extract", "2027-01-13",
     "notice lines with value dates from 1 April 2016 to 31 December 2026; the administrator re-keyed the former administrator's notices from 1 April 2016 when it took over the funds"),
    (C.F_LEGACY, "Former fund administrator, notice history", "2016-07-08", "notices with payment dates from 3 January 2006 to 30 June 2016"),
    (C.F_FX, "Custodian, reference rates", "2027-01-04", "business days from 2 January 2006 to 31 December 2026"),
    (C.F_POLICY, "Board of directors", "2025-06-16", "policy adopted 12 June 2025"),
    (C.F_CAPACC, "Current fund administrator, year-end capital account summary", "2027-01-13", "vehicles with a position at 31 December 2026"),
    (C.F_POS, "Custodian, positions report", "2027-01-04", "accounts 6011-2207 and 6011-2215 at the close of 31 December 2026"),
    (C.F_CUSTACT, "Custodian, settled activity report", "2027-01-04", "account 6011-2207, transactions settled from 1 January to 31 December 2026, with the opening position"),
    (C.F_DICT, "Investment office", "2027-01-14", "field definitions for the files in this set"),
    (C.F_GRANTS, "Grants office", "2026-12-11", "the 2027 grants budget as approved by the board on 10 December 2026"),
    (C.F_PERF, "General partners' quarterly reports, compiled by the investment office", "2026-12-18", "funds reporting at 30 September 2026"),
]


def write_manifest(out, w):
    return C.write_csv(os.path.join(out, C.F_MANIFEST), ("file", "source", "extract_date", "coverage"), MANIFEST)


def build_pack(out, world=None):
    os.makedirs(out, exist_ok=True)
    w = world or World()
    n = {}
    n[C.F_REGISTER] = write_register(out, w)
    n[C.F_ACTIVITY] = write_activity(out, w)
    n[C.F_LEGACY] = write_legacy(out, w)
    n[C.F_FX] = write_fx(out, w)
    n[C.F_CAPACC] = write_capital_accounts(out, w)
    n[C.F_CUSTACT], n[C.F_POS] = write_custody(out, w)
    n[C.F_PERF] = write_perf(out, w)
    n[C.F_DICT] = write_dictionary(out, w)
    n[C.F_MANIFEST] = write_manifest(out, w)
    D.write_policy(out)
    D.write_grants(out)
    return n


if __name__ == "__main__":
    print(build_pack(sys.argv[1]))
