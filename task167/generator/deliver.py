"""task167 deliverables: the golden workbook and memo, figures.json and submission.md, every figure computed from the
shipped files by calc.py.

    python3 deliver.py TARGET_DIR GOLDEN_DIR FIGURES_JSON SUBMISSION_MD"""
from __future__ import annotations

import datetime as dt
import io
import json
import math
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import matplotlib                      # noqa: E402
matplotlib.use("Agg")
import matplotlib.pyplot as plt        # noqa: E402
from openpyxl import Workbook          # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side   # noqa: E402

import calc as K                       # noqa: E402
import common as C                     # noqa: E402
import cracker as CR                   # noqa: E402
import docs as D                       # noqa: E402
import postprocess as PP               # noqa: E402

STRATS = ("Buyout", "Growth equity", "Venture capital", "Private credit")
F_BOOK = "call_reserve_2027.xlsx"
F_MEMO = "call_reserve_memo_2027.pdf"
MEMO_DATE = dt.date(2027, 1, 15)
BOARD_DATE = dt.date(2027, 1, 21)       # the board's January meeting


WORDS = ("zero one two three four five six seven eight nine ten eleven twelve thirteen fourteen fifteen sixteen "
         "seventeen eighteen nineteen twenty").split()


def words(n):
    return WORDS[n] if 0 <= n < len(WORDS) else "%d" % n


def r25(x_musd):
    return int(math.floor(x_musd / 25.0 + 0.5) * 25)


def figures(target):
    p, book, g = K.main_call(target)
    fc = K.golden_call(book)
    nav = CR.nav(p, target)
    assets = CR.assets(target)
    q = K.qualifying(target)
    ca = K.read_capital_accounts(target)
    hist = [f for f in book.funds.values() if not f["programme"]]
    prog = [f for f in book.funds.values() if f["programme"]]
    line_hist = [f for f in hist if K.is_line(f)]
    line_prog = [f for f in prog if K.is_line(f)]
    cb = K.curve(book, lambda f: K.hist(f) and K.is_line(f))
    ca_ = K.curve(book, lambda f: K.hist(f) and not K.is_line(f))
    y2 = lambda c: sum(c[12:24])
    # the rivals block 1 names: one schedule pooled over every earlier fund, as rates on the undrawn balance and as
    # shares of commitment
    pool_rate = K.curve(book, K.hist, rate=True)
    pool_share = K.curve(book, K.hist)
    rival_rate = K.reserve(sum(K.forecast(book, lambda f: pool_rate, rate=True).values()))
    rival_share = K.reserve(sum(K.forecast(book, lambda f: pool_share).values()))
    # the first-year split: no fund between these two whole percentages, read with or without annualising
    raw = lambda f: sum(f["h"][:min(12, f["obs"])]) / f["commitment"]
    seen = [f for f in book.funds.values() if f["obs"] >= 9]
    gap_lo = math.ceil(100 * max(max(f["y1"], raw(f)) for f in seen if K.is_line(f)))
    gap_hi = math.floor(100 * min(min(f["y1"], raw(f)) for f in seen if not K.is_line(f)))
    bt = backtest(book)
    # the own schedule's total miss on the left-out vintages, commitment-weighted and equally weighted: the bound the
    # write-up states holds under both
    own_err = max(abs(sum(r[4] - r[2] for r in b_)) / sum(r[2] for r in b_) for b_ in (bt, backtest(book, "equal")))
    fig = dict(
        reserve_musd=g["reserve"],
        expected_calls_total_musd=r25(g["total"] / 1e6),
        expected_calls_total_unrounded_musd=round(g["total"] / 1e6, 1),
        expected_calls_by_strategy_musd={s: r25(g["by_strategy"][s] / 1e6) for s in STRATS},
        expected_calls_by_strategy_unrounded_musd={s: round(g["by_strategy"][s] / 1e6, 1) for s in STRATS},
        programme_funds=len(prog), programme_line_funds=len(line_prog),
        history_funds=len(hist), history_line_funds=len(line_hist),
        first_year_gap_pct=[gap_lo, gap_hi],
        second_year_share_ratio=round(y2(cb) / y2(ca_), 2),
        backtest_vintages=len(bt), backtest_pooled_low_vintages=sum(1 for r in bt if r[3] < r[2]),
        backtest_own_within_pct=math.ceil(100 * own_err),
        rival_reserves_musd=dict(pooled_rates=rival_rate, pooled_shares=rival_share),
        reserve_step_musd=K.RESERVE_STEP_MUSD, reserve_year=K.YEAR, statement_year=K.YEAR - 1,
        policy_sections=D.section_numbers(),
        nav_musd=nav, assets_musd=assets,
        eur_rate=p.ye_rate,
        transfer_musd=round(g["reserve"] - assets["total"], 1),
        programme_commitment_musd=round(sum(f["commitment"] * book.fx_of(f) for f in prog) / 1e6, 1),
    )
    return p, book, fc, fig, q, ca


# ------------------------------------------------------------------------------------------------ the workbook
HEAD = PatternFill("solid", fgColor="DCE6F1")
THIN = Side(style="thin", color="9AA5B1")


def _title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=13)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, size=9, color="555555")


def _table(ws, r0, head, rows, fmt=None, widths=None):
    for j, h in enumerate(head):
        c = ws.cell(row=r0, column=j + 1, value=h)
        c.font = Font(bold=True)
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = Border(bottom=THIN)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = ws.cell(row=r0 + 1 + i, column=j + 1, value=v)
            if fmt and j in fmt and isinstance(v, (int, float)):
                c.number_format = fmt[j]
            if isinstance(v, dt.date):
                c.number_format = "yyyy-mm-dd"
    if widths:
        for j, w in enumerate(widths):
            ws.column_dimensions[chr(65 + j)].width = w
    return r0 + 1 + len(rows)


def write_book(path, p, book, fc, fig, ca):
    wb = Workbook()
    ws = wb.active
    ws.title = "Reserve"
    _title(ws, "Call reserve for 2027", "Investment office, prepared %s for the board's January meeting. $ million." %
           MEMO_DATE.strftime("%-d %B %Y"))
    rows = [("Call reserve for 2027 (policy 2.2, rounded up to the next $50 million)", fig["reserve_musd"]),
            ("Capital calls expected in 2027, all funds", fig["expected_calls_total_musd"])]
    for s in STRATS:
        rows.append(("Capital calls expected in 2027, %s" % s.lower(), fig["expected_calls_by_strategy_musd"][s]))
    r = _table(ws, 4, ("Item", "$ million"), rows, fmt={1: "#,##0"}, widths=(70, 14))
    for i, (lab, v) in enumerate((("Qualifying assets in reserve account 6011-2207 at 31 December 2026 (Reserve assets)",
                                   fig["assets_musd"]["total"]),
                                  ("Transfer from the endowment pool to meet the reserve (policy 5.1)", fig["transfer_musd"]))):
        ws.cell(row=r + i, column=1, value=lab)
        c = ws.cell(row=r + i, column=2, value=v)
        c.number_format = "#,##0.0"
    r += 2
    notes = [
        "Expected calls are to the nearest $25 million and cover the funds committed before 1 January 2027, fees and "
        "expenses included, before any distribution set against a call (policy 2.3).",
        "Euro funds at %.4f, the custodian's reference rate for 31 December 2026 (policy 2.5)." % fig["eur_rate"],
        "Each fund is projected by its age in months from first close (policy 2.4) on the schedule of the funds whose first "
        "year was financed like its own (sheet Call schedules); fund detail on sheet Funds.",
    ]
    for i, t in enumerate(notes):
        ws.cell(row=r + 1 + i, column=1, value=t).font = Font(size=9, color="444444")

    ws2 = wb.create_sheet("Call schedules")
    _title(ws2, "Call schedules by year of age", "Share of commitment called in each year from first close, funds of the "
           "2006 to 2022 vintages observed through that year. Per cent.")
    cb = K.curve(book, lambda f: K.hist(f) and K.is_line(f), step=12)
    ca_ = K.curve(book, lambda f: K.hist(f) and not K.is_line(f), step=12)
    cp = K.curve(book, K.hist, step=12)
    nb = lambda sel, y: sum(1 for f in book.funds.values() if sel(f) and f["obs"] >= 12 * (y + 1))
    rows = []
    for y in range(10):
        rows.append((y + 1, round(100 * cb[y], 1), nb(lambda f: K.hist(f) and K.is_line(f), y), round(100 * ca_[y], 1),
                     nb(lambda f: K.hist(f) and not K.is_line(f), y), round(100 * cp[y], 1)))
    r = _table(ws2, 4, ("Year of age", "First year fees and expenses only", "Funds", "Other funds", "Funds", "All funds"),
               rows, fmt={1: "0.0", 3: "0.0", 5: "0.0"}, widths=(12, 16, 9, 14, 9, 12))
    ws2.cell(row=r + 1, column=1, value="Second-year calls, left out one vintage at a time: the funds whose first year was "
             "fees and expenses only, $ million").font = Font(bold=True)
    bt = backtest(book)
    r = _table(ws2, r + 2, ("Vintage", "Funds", "Called in year two", "All-fund schedule", "Own schedule"),
               [(v, n, round(a, 1), round(pl, 1), round(ow, 1)) for v, n, a, pl, ow in bt],
               fmt={2: "#,##0.0", 3: "#,##0.0", 4: "#,##0.0"})

    ws3 = wb.create_sheet("Funds")
    _title(ws3, "Funds active at 31 December 2026", "First-year calls as a share of commitment; expected 2027 calls in "
           "$ million at the year-end rate.")
    rows = []
    for fid, f in sorted(book.funds.items(), key=lambda kv: (kv[1]["vintage"], kv[1]["first_close"], kv[0])):
        if f["status"] != "Active":
            continue
        name = p.veh[f["codes"][0]]["Fund"].replace(", L.P.", "")
        rows.append((fid, name, f["strategy"], f["vintage"], f["first_close"], f["ccy"],
                     round(f["commitment"] / 1e6, 2), round(100 * f["y1"], 1) if f["y1"] is not None else None,
                     "fees and expenses only" if K.is_line(f) else "investments", round(fc[fid] / 1e6, 1)))
    end = _table(ws3, 4, ("Fund ID", "Fund", "Strategy", "Vintage", "First close", "Currency", "Commitment (m)",
                          "First-year calls %", "First year", "Expected 2027 calls ($m)"), rows,
                 fmt={6: "#,##0.00", 7: "0.0", 9: "#,##0.0"}, widths=(10, 36, 15, 8, 11, 9, 14, 12, 20, 14))
    ws3.cell(row=end, column=1, value="All funds").font = Font(bold=True)
    c = ws3.cell(row=end, column=10, value=fig["expected_calls_total_unrounded_musd"])
    c.number_format = "#,##0.0"
    c.font = Font(bold=True)
    c.border = Border(top=THIN)
    ws3.cell(row=end + 2, column=1, value="Commitment in the fund's currency. First-year calls are the calls of the first "
             "twelve months from first close, annualised for the funds that first closed after January 2026; "
             "the total is of the unrounded fund figures.").font = Font(size=9, color="444444")
    ws3.freeze_panes = "C5"
    ws3.auto_filter.ref = "A4:J%d" % (end - 1)

    ws4 = wb.create_sheet("NAV")
    _title(ws4, "Private-markets net asset value at 31 December 2026", "Policy 4.1: each fund's latest reported value "
           "adjusted for the contributions paid and distributions received after its date, to 31 December; euros at %.4f. "
           "$ million." % fig["eur_rate"])
    nav = fig["nav_musd"]
    rows = [(s, nav["strategy:" + s]) for s in STRATS]
    rows += [("Euro funds", nav["ccy:EUR"]), ("Dollar funds", nav["ccy:USD"]), ("Vintages 2015 to 2018", nav["band:2015-2018"]),
             ("Vintages 2019 to 2022", nav["band:2019-2022"]), ("Total", nav["total"])]
    r = _table(ws4, 4, ("Net asset value", "$ million"), rows, fmt={1: "#,##0.0"}, widths=(34, 14))
    ws4.cell(row=r + 1, column=1, value="By fund (main and parallel vehicles together; the administrator's main-vehicle "
             "statements cover both)").font = Font(bold=True)
    det = nav_detail(p, ca)
    end = _table(ws4, r + 2, ("Fund ID", "Fund", "Strategy", "Vintage", "Currency", "Reported NAV (m)", "NAV date",
                              "Contributions after (m)", "Distributions after (m)", "Value 31 Dec 2026 ($m)"), det,
                 fmt={5: "#,##0.00", 7: "#,##0.00", 8: "#,##0.00", 9: "#,##0.0"})
    ws4.cell(row=end, column=1, value="All funds").font = Font(bold=True)
    c = ws4.cell(row=end, column=10, value=nav["total"])
    c.number_format = "#,##0.0"
    c.font = Font(bold=True)
    c.border = Border(top=THIN)
    ws4.cell(row=end + 2, column=1, value="Reported NAV, contributions and distributions in the fund's currency, from the "
             "administrator's capital account summary and the notices after each NAV date; a reported NAV below zero is the "
             "general partner's figure for a fund whose first deals sit on its credit line.").font = Font(size=9, color="444444")
    for col, w in zip("CDEFGHIJ", (15, 8, 9, 14, 11, 14, 14, 14)):
        ws4.column_dimensions[col].width = w

    ws5 = wb.create_sheet("Reserve assets")
    _title(ws5, "Reserve account 6011-2207, qualifying assets at 31 December 2026", "Policy 3.1: Treasury bills maturing "
           "by 31 December 2027, the Schedule 1 fund and cash, at market value on a settled basis. $ million.")
    a = fig["assets_musd"]
    rows = [("Qualifying assets", a["total"]), ("of which Treasury bills", a["bills"]),
            ("Bills maturing January to March 2027, face value", a["q1"]),
            ("Bills maturing April to June 2027, face value", a["q2"]),
            ("Bills maturing July to September 2027, face value", a["q3"]),
            ("Bills maturing October to December 2027, face value", a["q4"])]
    r = _table(ws5, 4, ("Item", "$ million"), rows, fmt={1: "#,##0.0"}, widths=(58, 14))
    rec = asset_reconciliation(os.path.join(os.path.dirname(os.path.dirname(path)), "target"))
    ws5.cell(row=r + 1, column=1, value="From the custodian's positions report").font = Font(bold=True)
    _table(ws5, r + 2, ("Line", "$ million"), rec, fmt={1: "#,##0.0"})
    wb.save(path)
    PP.normalize_ooxml(path, "Troy Stein", "Troy Stein", "2027-01-14T15:20:00Z", "2027-01-15T22:41:00Z", excel=True,
                       company=C.ORG)


def backtest(book, weight="commitment"):
    """second-year calls of the line-bridged funds, each vintage left out of the fit in turn"""
    out = []
    line = [f for f in book.funds.values() if K.hist(f) and K.is_line(f) and f["obs"] >= 24]
    for v in sorted(set(f["vintage"] for f in line)):
        sel_out = (lambda f, v=v: K.hist(f) and f["vintage"] != v)
        cpv = K.curve(book, sel_out, weight=weight)
        cbv = K.curve(book, lambda f, v=v: sel_out(f) and K.is_line(f), weight=weight)
        g = [f for f in line if f["vintage"] == v]
        a = sum(sum(f["h"][12:24]) for f in g)
        c = sum(f["commitment"] for f in g)
        out.append((v, len(g), a / 1e6, sum(cpv[12:24]) * c / 1e6, sum(cbv[12:24]) * c / 1e6))
    return out


def nav_detail(p, ca):
    out = []
    for d in ca:
        v = p.veh[d["Fund code"]]
        if v["Vehicle"] == "Parallel":
            continue
        codes = [c for c, x in p.veh.items() if x["fund"] == v["fund"]]
        nd = d["NAV date"]
        start = K.mi(nd) + 1 if nd else -1
        con = sum(a for c in codes for m, a in p.calls[c].items() if start <= m <= K.mi(K.YE))
        dis = sum(a for c in codes for m, a in p.dists[c].items() if start <= m <= K.mi(K.YE))
        nav0 = d["Reported NAV"] or 0.0
        rate = p.ye_rate if v["Currency"] == "EUR" else 1.0
        out.append((v["fund"], v["Fund"].replace(", L.P.", ""), v["Strategy"], v["Vintage"], v["Currency"],
                    round(nav0 / 1e6, 2), nd, round(con / 1e6, 2), round(dis / 1e6, 2), round((nav0 + con - dis) * rate / 1e6, 1)))
    return out


def asset_reconciliation(target):
    pos = K.read_positions(target)
    held = K.settled_holdings(target)
    res = [r for r in pos if r["account"] == "6011-2207"]
    tot = sum(float(r["market_value_usd"]) for r in res)
    unsettled = sum(float(r["market_value_usd"]) for r in res if r["asset_type"] == "Government Bill" and r["asset_id"] not in held)
    agency = sum(float(r["market_value_usd"]) for r in res if r["asset_type"] == "Government Bill" and r["asset_id"] in held
                 and not r["description"].startswith("U.S. TREASURY BILL"))
    note = sum(float(r["market_value_usd"]) for r in res if r["asset_type"] == "Government Note")
    prime = sum(float(r["market_value_usd"]) for r in res if r["asset_type"] == "Money Market Fund"
                and "Government Money Market Fund" not in r["description"])
    q = tot - unsettled - agency - note - prime
    m = lambda x: round(x / 1e6, 1)
    return [("Account total in the positions report", m(tot)),
            ("less bills bought 31 December for settlement 4 January 2027", -m(unsettled)),
            ("less agency discount notes in the custodian's Government Bill class", -m(agency)),
            ("less the Treasury note maturing 31 January 2028", -m(note)),
            ("less the prime money market fund (not in Schedule 1)", -m(prime)),
            ("Qualifying assets", m(q))]


# ------------------------------------------------------------------------------------------------- the memo
def chart_png(fig):
    vals = [fig["expected_calls_by_strategy_unrounded_musd"][s] for s in STRATS]
    lab = [fig["expected_calls_by_strategy_musd"][s] for s in STRATS]
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 8.5, "svg.hashsalt": "167"})
    f, ax = plt.subplots(figsize=(6.4, 3.1), dpi=200)
    base = 0.0
    colors = ("#2b3a55", "#4f6d8f", "#8aa4c2", "#c3d1e2")
    names = ("Buyout", "Growth equity", "Venture capital", "Private credit")
    for i, (v, l) in enumerate(zip(vals, lab)):
        ax.bar(i, v, bottom=base, color=colors[i], width=0.62, edgecolor="white")
        ax.text(i, base + v / 2, "%d" % l, ha="center", va="center", color="white" if i < 2 else "#1d1f22", fontsize=8.5)
        base += v
    ax.bar(4, base, color="#5b6068", width=0.62)
    ax.text(4, base / 2, "%d" % fig["expected_calls_total_musd"], ha="center", va="center", color="white", fontsize=9)
    ax.axhline(fig["reserve_musd"], color="#b2462e", lw=1.4, ls="--")
    ax.text(4.42, fig["reserve_musd"] + 6, "Reserve %d" % fig["reserve_musd"], ha="right", va="bottom", color="#b2462e",
            fontsize=8.5)
    ax.set_xticks(range(5))
    ax.set_xticklabels(list(names) + ["All funds"])
    ax.set_ylabel("$ million")
    ax.set_ylim(0, fig["reserve_musd"] * 1.12)
    ax.spines[["top", "right"]].set_visible(False)
    ax.set_title("Expected 2027 calls of \\$%d million by strategy, under the \\$%d million reserve"
                 % (fig["expected_calls_total_musd"], fig["reserve_musd"]), fontsize=9.5, loc="left")
    f.tight_layout()
    buf = io.BytesIO()
    f.savefig(buf, format="png", metadata={"Software": None})
    plt.close(f)
    return buf.getvalue()


def write_memo(path, fig):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    st = D._styles()
    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=24 * mm, rightMargin=24 * mm, topMargin=20 * mm,
                            bottomMargin=18 * mm, title="Call reserve for 2027", author=C.ORG, subject="", creator="",
                            producer="")
    e = fig["expected_calls_by_strategy_musd"]
    a = fig["assets_musd"]
    face = round(a["q1"] + a["q2"] + a["q3"] + a["q4"], 1)
    story = [Paragraph("Brandvold Foundation, investment office", st["sub"]),
             Paragraph("Call reserve for 2027: $%d million" % fig["reserve_musd"], st["h1"]),
             Paragraph("To: Board of directors, meeting of %s. From: Troy Stein, director of investment risk. Date: %s."
                       % (BOARD_DATE.strftime("%-d %B %Y"), MEMO_DATE.strftime("%-d %B %Y")), st["sub"])]
    paras = [
        "<b>The call reserve for 2027 is $%d million.</b> Under section 5.1 of the liquidity policy I certify it, with "
        "the reserve account's qualifying assets of $%.1f million at 31 December 2026; the shortfall of $%.1f million "
        "is to be transferred from the endowment pool within ten business days of the meeting."
        % (fig["reserve_musd"], a["total"], fig["transfer_musd"]),
        "We expect the funds committed before 1 January to call $%d million in 2027: buyout $%d million, growth equity "
        "$%d million, venture capital $%d million and private credit $%d million (to the nearest $25 million). Most of it "
        "is the 2026 programme in its second year. %s of its %s funds called only fees and expenses in 2026, as %d of our "
        "%d earlier funds did in their first year; those funds then called more than twice the share of commitment the "
        "others called in their second year. One schedule fitted over every fund, the usual pacing model, puts that "
        "second year low in each of the %d vintages we can test, and would set the reserve at $%d million."
        % (fig["expected_calls_total_musd"], e["Buyout"], e["Growth equity"], e["Venture capital"], e["Private credit"],
           words(fig["programme_line_funds"]).capitalize(), words(fig["programme_funds"]), fig["history_line_funds"],
           fig["history_funds"], fig["backtest_vintages"], fig["rival_reserves_musd"]["pooled_rates"]),
        "The reserve account held $%.1f million of qualifying assets at 31 December, $%.1f million of them Treasury bills "
        "maturing in 2027. The bills bought on 31 December for the 2027 ladder settle on 4 January and count from then; "
        "the agency discount notes the custodian reports with its bills, the 2028 note and the prime fund do not qualify "
        "(section 3.1 and Schedule 1)." % (a["total"], a["bills"]),
    ]
    for t in paras:
        story.append(Paragraph(t, st["body"]))
    rows = [["Treasury bills maturing in 2027, face value, $ million", "", "", "", ""],
            ["January to March", "April to June", "July to September", "October to December", "Total"],
            ["%.1f" % a["q1"], "%.1f" % a["q2"], "%.1f" % a["q3"], "%.1f" % a["q4"], "%.1f" % face]]
    tb = Table(rows, colWidths=[32 * mm] * 5)
    tb.setStyle(TableStyle([
        ("SPAN", (0, 0), (-1, 0)), ("FONT", (0, 0), (-1, 0), "LibSerif-Bold", 9.4),
        ("FONT", (0, 1), (-1, 1), "Carlito", 8.8), ("FONT", (0, 2), (-1, 2), "LibSerif", 9.6),
        ("ALIGN", (0, 1), (-1, -1), "RIGHT"), ("TEXTCOLOR", (0, 0), (-1, -1), D.INK),
        ("LINEBELOW", (0, 1), (-1, 1), 0.5, colors.HexColor("#9AA5B1")),
        ("LINEABOVE", (0, 0), (-1, 0), 0.5, colors.HexColor("#9AA5B1")),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5), ("TOPPADDING", (0, 0), (-1, -1), 2.5)]))
    story.append(Spacer(1, 2))
    story.append(tb)
    story.append(Paragraph("Source: custodian's positions report for account 6011-2207 at 31 December 2026 and its 2026 "
                           "settled activity.", st["small"]))
    story.append(Spacer(1, 6))
    story.append(Image(io.BytesIO(chart_png(fig)), width=160 * mm, height=77.5 * mm))
    story.append(Paragraph("Expected calls to the nearest $25 million on the commitments made before 1 January 2027, fees "
                           "and expenses included and before any distribution set against a call; euro funds at %.4f, "
                           "the custodian's reference rate for 31 December 2026 (sections 2.2 to 2.5). Fund detail and the "
                           "schedules: call_reserve_2027.xlsx." % fig["eur_rate"], st["small"]))
    doc.build(story, onFirstPage=D._footer("Brandvold Foundation, call reserve 2027"),
              onLaterPages=D._footer("Brandvold Foundation, call reserve 2027"))
    PP.normalize_pdf(path, D.pdf_stamp(MEMO_DATE, 16, 30), D.pdf_stamp(MEMO_DATE, 16, 30))


# ------------------------------------------------------------------------------------------------ submission.md
def submission(fig):
    e = fig["expected_calls_by_strategy_musd"]
    nav = fig["nav_musd"]
    a = fig["assets_musd"]
    L = []
    L.append("## Tags")
    L.append("")
    L.append("**Domain:** General Statistics & Data Science (regression, classification and model diagnostics: a grant-making foundation's "
             "2027 call reserve for its private fund commitments, each fund's call schedule read from the funds whose first "
             "year was financed like its own and validated on the funds that have passed their second year).")
    L.append("**Analytical objective:** Forecasting & Predictive Modeling.")
    L.append("")
    L.append("## 1. Final Recommendation")
    L.append("")
    rv = fig["rival_reserves_musd"]
    L.append("**The Foundation should hold a call reserve of $%d million for 2027.** The funds committed before 1 January "
             "are expected to call $%d million in 2027, most of it the 2026 programme in its second year: %s of its "
             "%s funds called only fees and expenses in 2026, as %d of the %d earlier funds did in their first year "
             "before calling more than twice the others' share of commitment in their second. Not $%d million, the reserve "
             "from Takahashi-Alexander rates on the undrawn balance pooled over every earlier fund ($%d million from shares "
             "of commitment pooled the same way), which puts the fee-only funds' second-year calls low in every vintage "
             "left out of the fit. Not 2026's calls carried forward, which take the programme's first year as its pace."
             % (fig["reserve_musd"], fig["expected_calls_total_musd"], words(fig["programme_line_funds"]),
                words(fig["programme_funds"]), fig["history_line_funds"], fig["history_funds"], rv["pooled_rates"],
                rv["pooled_shares"]))
    L.append("")
    L.append("## 2. Critical Components")
    L.append("")
    L.append("1. Funds whose first year called only fees and expenses: **%d** of the %d earlier funds and **%d** of the "
             "%d programme funds." % (fig["history_line_funds"], fig["history_funds"], fig["programme_line_funds"],
                                     fig["programme_funds"]))
    L.append("2. Capital calls expected in 2027: **$%d million** (to the nearest $25 million)." % fig["expected_calls_total_musd"])
    L.append("3. By strategy: buyout **$%d million**, growth equity **$%d million**, venture capital **$%d million**, private "
             "credit **$%d million**." % (e["Buyout"], e["Growth equity"], e["Venture capital"], e["Private credit"]))
    L.append("")
    L.append("## 3. Step-by-Step Solution")
    L.append("")
    L.append("1. Joined investor_capital_activity_2016-04_2026-12.csv and legacy_administrator_notices_2006-2016.xlsx to "
             "fund_commitments_register.xlsx (the legacy file before 1 April 2016, settled notices only), and built each "
             "fund's calls gross of distributions (policy 2.3) by month of age from its first close (policy 2.4).")
    L.append("2. Split the funds on their first twelve months from first close: %d earlier funds and %d programme funds "
             "called only fees and expenses, and no fund called between %d and %d per cent of commitment."
             % (fig["history_line_funds"], fig["programme_line_funds"], fig["first_year_gap_pct"][0], fig["first_year_gap_pct"][1]))
    L.append("3. Left out one vintage at a time, the schedule pooled over every earlier fund put the fee-only funds' "
             "second-year calls low in all %d vintages, and their own schedule came within %d per cent of those calls in "
             "total." % (fig["backtest_vintages"], fig["backtest_own_within_pct"]))
    L.append("4. Fitted each class's schedule by month of age on the earlier funds, the fee-only funds calling more than twice "
             "the others' share of commitment in their second year.")
    L.append("5. Applied each fund's class schedule to its 2027 ages, capped at its undrawn commitment, euro funds at %.4f "
             "(policy 2.5): $%d million expected." % (fig["eur_rate"], fig["expected_calls_total_musd"]))
    L.append("6. Rounded up to the next $50 million (policy 2.2): a call reserve of $%d million." % fig["reserve_musd"])
    L.append("")
    L.append("## 4. Deliverable Answers")
    L.append("")
    L.append("### call_reserve_2027.xlsx")
    L.append("")
    L.append("1. Call reserve for 2027: $%d million." % fig["reserve_musd"])
    L.append("2. Capital calls expected in 2027, $ million to the nearest 25: all funds %d; buyout %d; growth equity %d; "
             "venture capital %d; private credit %d." % (fig["expected_calls_total_musd"], e["Buyout"], e["Growth equity"],
                                                        e["Venture capital"], e["Private credit"]))
    L.append("3. Private-markets net asset value at 31 December 2026, $ million:")
    for s in STRATS:
        L.append("   - %s: %.1f" % (s, nav["strategy:" + s]))
    L.append("   - Euro funds: %.1f" % nav["ccy:EUR"])
    L.append("   - Dollar funds: %.1f" % nav["ccy:USD"])
    L.append("   - Vintages 2015 to 2018: %.1f" % nav["band:2015-2018"])
    L.append("   - Vintages 2019 to 2022: %.1f" % nav["band:2019-2022"])
    L.append("   - Total: %.1f" % nav["total"])
    L.append("4. Reserve account at 31 December 2026, $ million:")
    L.append("   - Qualifying assets: %.1f" % a["total"])
    L.append("   - Treasury bills among them: %.1f" % a["bills"])
    L.append("   - Face value maturing January to March 2027: %.1f" % a["q1"])
    L.append("   - Face value maturing April to June 2027: %.1f" % a["q2"])
    L.append("   - Face value maturing July to September 2027: %.1f" % a["q3"])
    L.append("   - Face value maturing October to December 2027: %.1f" % a["q4"])
    L.append("")
    L.append("### call_reserve_memo_2027.pdf")
    L.append("")
    L.append("1. States the call reserve of $%d million." % fig["reserve_musd"])
    L.append("2. Chart of the capital calls expected in 2027 by strategy, stacked to the $%d million total, with the $%d "
             "million reserve drawn across as a line." % (fig["expected_calls_total_musd"], fig["reserve_musd"]))
    return "\n".join(L) + "\n"


def main(target, golden, fig_path, sub_path):
    os.makedirs(golden, exist_ok=True)
    p, book, fc, fig, q, ca = figures(target)
    write_book(os.path.join(golden, F_BOOK), p, book, fc, fig, ca)
    write_memo(os.path.join(golden, F_MEMO), fig)
    flat = dict(fig)
    with open(fig_path, "w") as fh:
        json.dump(flat, fh, indent=1, sort_keys=True)
        fh.write("\n")
    with open(sub_path, "w") as fh:
        fh.write(submission(fig))
    return fig


if __name__ == "__main__":
    main(*sys.argv[1:5])
