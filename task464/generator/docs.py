"""The pack's documents: the 2027 grant rules, the retiring test's interpretive guide, the new test's grade 8 technical
summary, the bridge study memo, the State Board's August cut-score resolution, the district directory workbook, the
folder's data dictionary and its file register."""
import datetime as dt

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import letter                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import inch                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import numpy as np    # noqa: E402

import pack as P      # noqa: E402
import world as W     # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("LiberationSerif", "Regular"), ("LiberationSerif-Bold", "Bold"), ("LiberationSans", "Regular"),
              ("LiberationSans-Bold", "Bold")):
    try:
        pdfmetrics.getFont(nm)
    except KeyError:
        pdfmetrics.registerFont(TTFont(nm, FONT_DIR + nm.split("-")[0] + "-" + f + ".ttf"))
INK = colors.HexColor("#1b1d1f")
GREY = colors.HexColor("#5b6067")
NAVY = colors.HexColor("#1f3a5f")
GREEN = colors.HexColor("#2f5d3a")
MAROON = colors.HexColor("#6b2737")
SLATE = colors.HexColor("#3d4f5c")
BODY = ParagraphStyle("body", fontName="LiberationSerif", fontSize=10.5, leading=14.2, textColor=INK, spaceAfter=6)
SANS = ParagraphStyle("sans", fontName="LiberationSans", fontSize=9.6, leading=13.2, textColor=INK, spaceAfter=6)
CELL = ParagraphStyle("cell", fontName="LiberationSans", fontSize=8.8, leading=11, textColor=INK)
CELLB = ParagraphStyle("cellb", parent=CELL, fontName="LiberationSans-Bold")
SMALL = ParagraphStyle("small", fontName="LiberationSans", fontSize=8.5, leading=11, textColor=GREY)


def styles(accent):
    h1 = ParagraphStyle("h1", fontName="LiberationSans-Bold", fontSize=14, leading=18, textColor=accent, spaceAfter=2)
    h1s = ParagraphStyle("h1s", fontName="LiberationSans", fontSize=9.3, leading=12.5, textColor=GREY, spaceAfter=10)
    h2 = ParagraphStyle("h2", fontName="LiberationSans-Bold", fontSize=10.5, leading=14, textColor=INK, spaceBefore=7,
                        spaceAfter=3)
    return h1, h1s, h2


def _doc(path, title, author, footer, layout="left"):
    """layout: left (footer text left, page right), centre (one centred line), header (a ruled running head and a
    centred page number), vendor (footer right with the page)."""
    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(GREY)
        if layout == "left":
            canvas.setFont("LiberationSans", 7.5)
            canvas.drawString(0.9 * inch, 0.6 * inch, footer)
            canvas.drawRightString(letter[0] - 0.9 * inch, 0.6 * inch, "Page %d" % doc.page)
        elif layout == "centre":
            canvas.setFont("LiberationSerif", 8.5)
            canvas.drawCentredString(letter[0] / 2, 0.6 * inch, "%s  -  %d" % (footer, doc.page))
        elif layout == "header":
            canvas.setFont("LiberationSans", 8)
            canvas.drawString(0.9 * inch, letter[1] - 0.55 * inch, footer)
            canvas.setStrokeColor(GREY)
            canvas.setLineWidth(0.5)
            canvas.line(0.9 * inch, letter[1] - 0.6 * inch, letter[0] - 0.9 * inch, letter[1] - 0.6 * inch)
            canvas.drawCentredString(letter[0] / 2, 0.55 * inch, str(doc.page))
        else:
            canvas.setFont("LiberationSans", 7)
            canvas.drawRightString(letter[0] - 0.9 * inch, 0.6 * inch, "%s  |  page %d" % (footer, doc.page))
        canvas.restoreState()
    d = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.9 * inch, rightMargin=0.9 * inch, topMargin=0.8 * inch,
                          bottomMargin=0.9 * inch, title=title, author=author, subject="", creator="", producer="")
    return d, frame


def _table(rows, widths, head="#e6eef0"):
    data = [[Paragraph(str(c), CELLB if i == 0 else CELL) for c in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa0a6")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(head)),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 4),
                           ("RIGHTPADDING", (0, 0), (-1, -1), 4)]))
    return t


def fmt(v):
    return "{:,}".format(int(v))


def long_date(iso):
    d = dt.date.fromisoformat(iso)
    return d.strftime("%B ") + str(d.day) + d.strftime(", %Y")


# ---------------------------------------------------------------- the 2027 grant rules
NOTICE_DAYS = {"1010": "2026-09-28", "1020": "2026-10-02", "1110": "2026-10-09", "1120": "2026-09-30",
               "1210": "2026-10-14", "1220": "2026-10-06", "1230": "2026-10-13", "1240": "2026-10-15",
               "1310": "2026-10-08", "1320": "2026-09-24", "1410": "2026-10-12", "1440": "2026-10-01",
               "1520": "2026-10-15", "1700": "2026-10-07"}


def write_rules(path):
    H1, H1S, H2 = styles(GREEN)
    doc, frame = _doc(path, "Adolescent Literacy Partnership: 2027 Grant Rules", W.LITERACY_OFFICE,
                      "%s, %s  |  Adolescent Literacy Partnership  |  2027 grant rules" % (W.STATE_AGENCY,
                                                                                         W.LITERACY_OFFICE))
    names = {c: n for c, n, *_ in W.DISTRICTS}
    s = [Paragraph("Adolescent Literacy Partnership: 2027 Grant Rules", H1),
         Paragraph("%s, %s. Issued September 8, 2026 by %s, %s. Attachment A added October 16, 2026." % (
             W.STATE_AGENCY, W.LITERACY_OFFICE, W.PEOPLE["director"], W.TITLES["director"]), H1S),
         Paragraph("1. Purpose", H2),
         Paragraph("The Adolescent Literacy Partnership pays for reading intervention for students in their first two "
                   "years of high school who are reading well below grade level. A grant funds reading "
                   "interventionists, their training and their materials in the funded district's high schools.", BODY),
         Paragraph("2. Grants", H2),
         Paragraph("The State Board of Education will award six grants of $410,000 each for the calendar year 2027. "
                   "The amount of a grant does not depend on the number of students counted.", BODY),
         Paragraph("3. Selection", H2),
         Paragraph("A district that filed a notice of intent with the Office of Literacy by %s is an applicant; the "
                   "applicants are listed in Attachment A. The Board will fund the six applicants with the most "
                   "eligible students at its meeting of December 10, 2026." % long_date(W.NOTICE_DEADLINE), BODY),
         Paragraph("4. Eligible students", H2),
         Paragraph("An eligible student is a student in grade 9 or grade 10 on the applicant's roll on %s whose "
                   "latest grade 8 state reading result is at Level 1 of the %s, that is, a grade 8 scale score below "
                   "%d. A result on the %s, which the Measure replaced after the spring 2025 administration, counts at "
                   "the level it places at on the Measure's scale. A student without a grade 8 state reading result "
                   "is not counted." % (long_date(W.COUNT_DAY), W.NEW_TEST, W.NEW_CUTS[8][0], W.OLD_TEST), BODY),
         Paragraph("5. Conditions", H2),
         Paragraph("A funded district names a lead interventionist by January 15, 2027, serves students in grades 9 "
                   "and 10 in every high school it operates, and reports students served and sessions held each "
                   "quarter. Unspent funds are returned to the Department after December 31, 2027.", BODY),
         PageBreak(),
         Paragraph("Attachment A. Applicants", H2),
         _table([["District code", "District", "Notice of intent received"]] +
                [[c, names[c], long_date(NOTICE_DAYS[c])] for c in W.APPLICANTS], [80, 250, 140], head="#e3ece5")]
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------- the retiring test's interpretive guide
def write_glra_guide(path):
    H1, H1S, H2 = styles(MAROON)
    doc, frame = _doc(path, "Grade-Level Reading Assessment, Spring 2025: Interpretive Guide", W.OLD_VENDOR,
                      "Grade-Level Reading Assessment  |  Spring 2025 Interpretive Guide  |  %s" % W.OLD_VENDOR,
                      layout="header")
    tables = {g: W.glra_table(g) for g in (6, 7, 8)}
    conv = [["Raw score", "Grade 6", "Grade 7", "Grade 8"]] + [[r, tables[6][r], tables[7][r], tables[8][r]]
                                                                for r in range(W.GLRA_ITEMS + 1)]
    lev = [["Grade", "Below Basic", "Basic", "Proficient", "Advanced"]]
    for g in (6, 7, 8):
        a, b, c = W.OLD_CUTS[g]
        lev.append([g, "%d to %d" % (W.OLD["lo"], max(v for v in tables[g] if v < a)),
                    "%d to %d" % (a, max(v for v in tables[g] if v < b)), "%d to %d" % (b, max(v for v in tables[g] if v < c)),
                    "%d to %d" % (c, W.OLD["hi"])])
    s = [Paragraph("Grade-Level Reading Assessment, Spring 2025", H1),
         Paragraph("Interpretive guide for district and school staff. Prepared by %s for the %s, %s. August 15, "
                   "2025." % (W.OLD_VENDOR, W.STATE_AGENCY, W.OFFICE), H1S),
         Paragraph("1. The assessment", H2),
         Paragraph("The Grade-Level Reading Assessment measures reading in grades 3 to 8. Each grade sits one fixed "
                   "paper form of %d multiple-choice items, each scored right or wrong. The spring 2025 window ran from "
                   "April 7 to April 25, 2025, with make-up sessions from May 5 to May 16. Spring 2025 is the last "
                   "administration: the Department replaces the assessment with the %s from spring 2026." % (
                       W.GLRA_ITEMS, W.NEW_TEST), SANS),
         Paragraph("2. Scale scores", H2),
         Paragraph("A student's raw score (items correct) converts to a scale score from %d to %d through the form's "
                   "conversion table (Table 1). The tables come from the forms' item calibrations; every raw score has "
                   "one scale score, so only the scale scores in a grade's table occur. Scale scores of different "
                   "grades are on one scale." % (W.OLD["lo"], W.OLD["hi"]), SANS),
         _table(conv[:1] + conv[1:], [70, 70, 70, 70], head="#efe2e5"),
         Paragraph("Table 1. Raw score to scale score, spring 2025 forms, grades 6 to 8.", SMALL),
         Spacer(1, 6),
         Paragraph("3. Performance levels", H2),
         Paragraph("Each scale score falls in one of four performance levels (Table 2). Students at Below Basic read "
                   "well below the grade's expectations; students at Proficient or Advanced meet or exceed them.", SANS),
         _table(lev, [50, 95, 95, 95, 95], head="#efe2e5"),
         Paragraph("Table 2. Scale score ranges of the performance levels, grades 6 to 8.", SMALL),
         Paragraph("4. Precision", H2),
         Paragraph("The forms' internal consistency (KR-20) in spring 2025 was 0.85 in grade 6, 0.86 in grade 7 and "
                   "0.87 in grade 8. The standard error of measurement is about 20 scale points near the middle of the "
                   "scale and larger toward its ends, where the table's steps are wider.", SANS),
         Paragraph("5. The student file", H2),
         Paragraph("Each answer booklet carries a Calderwood test-taker number. Booklets printed for the Department's "
                   "registration carry a pre-ID label with the student's state identifier; booklets a school prepared "
                   "on the day carry the test-taker number only, and Calderwood matches them to state identifiers "
                   "after scoring. SCORE_STATUS is SCR (scored), ABS (absent), EXM (exempt) or INV (invalidated). A "
                   "make-up session after an invalidation is a further record under the same test-taker number with "
                   "SESSION Make-up.", SANS)]
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------- the new test's grade 8 technical summary
def slm_grade8_stats(x):
    sc = []
    for s, r, school, dist in P.slm_records(x):
        if r["grade"] != 8:
            continue
        v = r["score"] if r["status"] == "scored" else (r["retake_score"] if r["status"] == "invalid" and r["retake"]
                                                       else None)
        if v is not None:
            sc.append(v)
    sc = np.array(sc)
    c = W.NEW_CUTS[8]
    shares = [np.mean(sc < c[0]), np.mean((sc >= c[0]) & (sc < c[1])), np.mean((sc >= c[1]) & (sc < c[2])),
              np.mean(sc >= c[2])]
    return dict(n=len(sc), mean=float(sc.mean()), sd=float(sc.std(ddof=1)), shares=[100 * v for v in shares])


def write_slm_summary(path, x):
    H1, H1S, H2 = styles(NAVY)
    st = slm_grade8_stats(x)
    doc, frame = _doc(path, "Statewide Literacy Measure, Spring 2026: Grade 8 Technical Summary", W.NEW_VENDOR,
                      "%s  |  SLM-TS-26-08" % W.NEW_VENDOR, layout="vendor")
    c = W.NEW_CUTS[8]
    s = [Paragraph("Statewide Literacy Measure, Spring 2026: Grade 8 Technical Summary", H1),
         Paragraph("Prepared by %s for the %s, %s. June 26, 2026." % (W.NEW_VENDOR, W.STATE_AGENCY, W.OFFICE), H1S),
         Paragraph("1. The Measure", H2),
         Paragraph("The %s is a computer-adaptive reading assessment for grades 3 to 8, first given in spring 2026. "
                   "Each student answers items chosen from a calibrated bank to suit the student's earlier answers, and "
                   "the score is the student's estimated reading ability, reported as a whole scale score from %d to "
                   "%d on one scale for all grades. The spring 2026 window ran from April 13 to May 15, 2026, with "
                   "make-up sessions from May 18 to May 29." % (W.NEW_TEST, W.NEW["lo"], W.NEW["hi"]), SANS),
         Paragraph("2. Grade 8 performance levels", H2),
         Paragraph("The State Board of Education adopted the grade 8 cut scores on %s (Resolution %s): Level 2 begins "
                   "at %d, Level 3 at %d and Level 4 at %d. Students at Level 1 read well below the expectations for "
                   "the end of grade 8; students at Level 3 meet them and students at Level 4 exceed them." % (
                       long_date(W.GRADE8_RESOLUTION[1]), W.GRADE8_RESOLUTION[0], c[0], c[1], c[2]), SANS),
         _table([["Level", "Grade 8 scale scores"], ["Level 1", "%d to %d" % (W.NEW["lo"], c[0] - 1)],
                 ["Level 2", "%d to %d" % (c[0], c[1] - 1)], ["Level 3", "%d to %d" % (c[1], c[2] - 1)],
                 ["Level 4", "%d to %d" % (c[2], W.NEW["hi"])]], [90, 140], head="#e1e8f0"),
         Paragraph("3. Precision", H2),
         Paragraph("The grade 8 test's marginal reliability in spring 2026 was 0.88. The standard error of a grade 8 "
                   "scale score is about 32 points across the range where most students score, including the Level 2 "
                   "cut.", SANS),
         Paragraph("4. Statewide grade 8 results, spring 2026", H2),
         _table([["Students with a score", "Mean scale score", "Standard deviation", "Level 1", "Level 2", "Level 3",
                  "Level 4"],
                 [fmt(st["n"]), "%.1f" % st["mean"], "%.1f" % st["sd"]] + ["%.1f%%" % v for v in st["shares"]]],
                [80, 62, 62, 52, 52, 52, 52], head="#e1e8f0"),
         Paragraph("Includes make-up sessions and students of nonpublic schools tested under agreement.", SMALL),
         Paragraph("5. The student file", H2),
         Paragraph("The student file holds one record for each student and session, with the testing district and "
                   "school, the session (Main or Make-up), its date, its status (Scored, Absent, Exempt or "
                   "Invalidated) and, for scored sessions, the scale score and its performance level.", SANS)]
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------- the bridge study memo
def bridge_stats(x):
    rows = x["bridge"]["rows"]
    old = np.array([r["old"] for r in rows], float)
    new = np.array([r["new"] for r in rows], float)
    return dict(n=len(rows), schools=len(x["bridge"]["schools"]), om=old.mean(), osd=old.std(ddof=1), nm=new.mean(),
                nsd=new.std(ddof=1), r=float(np.corrcoef(old, new)[0, 1]),
                first=sum(1 for r in rows if r["order"] == "GLRA"))


def write_bridge_memo(path, x):
    H1, H1S, H2 = styles(SLATE)
    st = bridge_stats(x)
    n_public = len(set(v for vs in x["mid"].values() for v in vs))
    doc, frame = _doc(path, "The spring 2026 bridge study of the two state reading tests", W.OFFICE,
                      "%s, %s  |  Memorandum  |  August 21, 2026" % (W.STATE_AGENCY, W.OFFICE))
    s = [Paragraph("The spring 2026 bridge study of the two state reading tests", H1),
         Paragraph("Memorandum. To: %s, %s. From: %s, %s. August 21, 2026." % (
             W.PEOPLE["director"], W.TITLES["director"], W.PEOPLE["psychometrician"], W.TITLES["psychometrician"]), H1S),
         Paragraph("1. Why the study was run", H2),
         Paragraph("With the spring 2025 administration the Department retired the %s. Its records now hold grade 8 "
                   "reading results on two tests with different scales: spring 2025 results on the retiring test's "
                   "scale of %d to %d and spring 2026 results on the %s's scale of %d to %d. The bridge study gives "
                   "the Department results on both tests from one group of students." % (
                       W.OLD_TEST, W.OLD["lo"], W.OLD["hi"], W.NEW_TEST, W.NEW["lo"], W.NEW["hi"]), SANS),
         Paragraph("2. Design", H2),
         Paragraph("We drew %d of the state's %d public schools with a grade 8 at random. In each, the grade 8 "
                   "students sat the retiring test's spring 2025 grade 8 form within a week of their %s session; "
                   "%s scored the booklets on the 2025 conversion table under its contract. Half of the schools, "
                   "assigned at random, sat the retiring test first and half sat it after the Measure." % (
                       st["schools"], n_public, W.NEW_TEST, W.OLD_VENDOR), SANS),
         Paragraph("3. Participation", H2),
         Paragraph("%s students have a scored result on both tests, %s of them in schools that sat the retiring test "
                   "first. Students who missed either session, or whose session was invalidated, are not in the file. "
                   "The Measure score in the file is the student's operational score." % (fmt(st["n"]),
                                                                                          fmt(st["first"])), SANS),
         _table([["", "Students", "Mean", "Standard deviation"],
                 ["%s (grade 8 form)" % W.OLD_TEST, fmt(st["n"]), "%.1f" % st["om"], "%.1f" % st["osd"]],
                 ["%s (grade 8)" % W.NEW_TEST, fmt(st["n"]), "%.1f" % st["nm"], "%.1f" % st["nsd"]]],
                [230, 70, 70, 100], head="#e2e7ea"),
         Paragraph("The correlation between the two scores is %.2f." % st["r"], SANS),
         Paragraph("4. Files", H2),
         Paragraph("The paired scores are in %s, one record for each student, with the order in which the school sat "
                   "the two tests." % W.FILES["bridge"], SANS)]
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------- the State Board's August resolution
LOWER_CUTS = {3: (2290, 2383, 2454), 4: (2315, 2408, 2479), 5: (2340, 2433, 2504)}


def write_resolution(path):
    H1, H1S, H2 = styles(NAVY)
    doc, frame = _doc(path, "Resolution %s: Statewide Literacy Measure cut scores, grades 3 to 7" % W.LOWER_RESOLUTION[0],
                      W.BOARD, "Resolution %s" % W.LOWER_RESOLUTION[0], layout="centre")
    cuts = dict(LOWER_CUTS)
    cuts.update({6: W.NEW_CUTS[6], 7: W.NEW_CUTS[7]})
    s = [Paragraph("Resolution %s" % W.LOWER_RESOLUTION[0], H1),
         Paragraph("%s. Performance level cut scores for the %s, grades 3 to 7. Adopted %s." % (
             W.BOARD, W.NEW_TEST, long_date(W.LOWER_RESOLUTION[1])), H1S),
         Paragraph("Whereas the Department first administered the %s in grades 3 to 8 in spring 2026; whereas "
                   "standard-setting panels of teachers met in June 2026 and the Department brought their work to the "
                   "Board with its recommendation; and whereas the Board adopted the grade 8 cut scores by Resolution "
                   "%s on %s;" % (W.NEW_TEST, W.GRADE8_RESOLUTION[0], long_date(W.GRADE8_RESOLUTION[1])), BODY),
         Paragraph("The Board resolves that the cut scores in Attachment A are the performance level cut scores of the "
                   "%s for grades 3 to 7, beginning with the spring 2026 results." % W.NEW_TEST, BODY),
         Paragraph("Attachment A. Scale score at which each level begins", H2),
         _table([["Grade", "Level 2", "Level 3", "Level 4"]] + [[g, cuts[g][0], cuts[g][1], cuts[g][2]]
                                                               for g in (3, 4, 5, 6, 7)], [60, 80, 80, 80],
                head="#e1e8f0"),
         Spacer(1, 8),
         Paragraph("Adopted by a vote of 9 to 0, one member absent.", BODY)]
    doc.build(s, onFirstPage=frame, onLaterPages=frame)


# ---------------------------------------------------------------- the district directory
HEADS = ["Ash", "Bel", "Cor", "Dun", "Ell", "Fair", "Gal", "Hart", "Ing", "Jess", "Kins", "Lath", "Merr", "Nor",
         "Pend", "Rad", "Sor", "Thack", "Whit", "Yard", "Bram", "Call", "Demp", "Ferr", "Gilch", "Hals", "Ives", "Kell",
         "Lang", "Morr", "Ness", "Orr", "Pres", "Quen", "Ross", "Sell", "Tall", "Verr", "Wess", "Alb", "Brin", "Carr",
         "Dals", "Edg", "Finn", "Hodd", "Kemp", "Linds", "Mack"]
TAILS = ["by", "more", "ton", "wood", "well", "field", "ley", "worth", "ford", "rick", "ham", "brook", "sey", "ridge",
         "mont", "stead", "ing", "dale", "ett", "cote"]
FEATURES = ["Cedar Ridge", "Willow Bend", "Stony Hollow", "Copper Hill", "Fox Meadow", "Birch Hollow", "Heron Marsh",
            "Juniper Hill", "Lark Valley", "Mill Pond", "Osprey Point", "Quarry Road", "Red Oak", "Sandstone",
            "Westwind", "Aspen Park", "Bluff View", "Deer Run", "Fern Hollow", "Granite Ledge"]
DIRECTIONS = ["North", "South", "East", "West", "Central"]
LOCALE = {"1010": "City", "1020": "City", "1110": "Suburb", "1120": "Suburb", "1210": "Town", "1220": "Suburb",
          "1230": "Town", "1240": "Rural"}


def _surnames():
    import itertools
    import re
    names = sorted(set(h + t for h, t in itertools.product(HEADS, TAILS)
                       if not re.search(r"(.)\1\1", (h + t).lower()) and not (h[-1] == t[0])))
    rng = np.random.default_rng(4641)
    return [names[i] for i in rng.permutation(len(names))]


def school_names(x):
    """Each public school's name: middle schools named for people and places as districts name them, high schools for
    the town (with a direction where a district runs several)."""
    pool = _surnames()
    feats = list(FEATURES)
    out = {}
    k_name = k_feat = 0
    for code, name, cty, *_ in W.DISTRICTS:
        subs = [s_ for s_, _, _ in W.MERGER[code]] if code in W.MERGER else [code]
        for sub in subs:
            for k, sc in enumerate(x["mid"][sub]):
                if (int(sc) % 5 == 0) and k_feat < len(feats):
                    out[sc] = "%s Middle School" % feats[k_feat]
                    k_feat += 1
                else:
                    out[sc] = "%s Middle School" % pool[k_name]
                    k_name += 1
        town = W.short_name(code).replace("-", " ") if code not in W.MERGER else "Ashgrove-Fenley"
        hs = x["high"][code]
        for k, sc in enumerate(hs):
            if k == 0:
                out[sc] = "%s High School" % town
            elif k <= len(DIRECTIONS):
                out[sc] = "%s %s High School" % (town, DIRECTIONS[k - 1])
            else:
                out[sc] = "%s High School" % pool[k_name]
                k_name += 1
    return out


def write_directory(path, x):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    names = school_names(x)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Districts"
    ws["A1"] = "School district directory, 2026-27"
    ws["A1"].font = Font(name="Calibri", size=13, bold=True)
    ws["A2"] = "%s, Office of Data Management. As of October 1, 2026." % W.STATE_AGENCY
    ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="595959")
    head = ["District code", "District", "County", "Locale", "Status"]
    for j, h in enumerate(head, start=1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(name="Calibri", size=10, bold=True)
        c.fill = PatternFill("solid", fgColor="DCE6F0")
    rows = []
    for code, name, cty, *_ in W.DISTRICTS:
        status = "Open"
        if code in W.MERGER:
            subs = W.MERGER[code]
            status = "Open; formed July 1, 2026 from %s (%s) and %s (%s)" % (subs[0][1], subs[0][0], subs[1][1],
                                                                             subs[1][0])
        rows.append([code, name, cty + " County", LOCALE.get(code, "Town" if int(code) % 3 else "Rural"), status])
    for new, parts in W.MERGER.items():
        cty = next(d[2] for d in W.DISTRICTS if d[0] == new)
        for sub, sname, _ in parts:
            rows.append([sub, sname, cty + " County", "Rural", "Closed June 30, 2026; consolidated into %s" % new])
    rows.sort()
    for i, r in enumerate(rows, start=5):
        for j, v in enumerate(r, start=1):
            ws.cell(row=i, column=j, value=v).font = Font(name="Calibri", size=10)
    for j, w in enumerate([13, 38, 18, 10, 70], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    ws2 = wb.create_sheet("Schools")
    ws2["A1"] = "Public schools, 2026-27"
    ws2["A1"].font = Font(name="Calibri", size=13, bold=True)
    ws2["A2"] = "Middle schools (grades 6 to 8) and high schools (grades 9 to 12), by district."
    ws2["A2"].font = Font(name="Calibri", size=9, italic=True, color="595959")
    head = ["School code", "School", "District code", "Level", "Grades"]
    for j, h in enumerate(head, start=1):
        c = ws2.cell(row=4, column=j, value=h)
        c.font = Font(name="Calibri", size=10, bold=True)
        c.fill = PatternFill("solid", fgColor="DCE6F0")
    srows = []
    for code, *_ in W.DISTRICTS:
        subs = [s for s, _, _ in W.MERGER[code]] if code in W.MERGER else [code]
        for sub in subs:
            for sc in x["mid"][sub]:
                srows.append([sc, names[sc], code, "Middle", "6-8"])
        for sc in x["high"][code]:
            srows.append([sc, names[sc], code, "High", "9-12"])
    srows.sort()
    for i, r in enumerate(srows, start=5):
        for j, v in enumerate(r, start=1):
            ws2.cell(row=i, column=j, value=v).font = Font(name="Calibri", size=10)
    for j, w in enumerate([12, 40, 13, 9, 8], start=1):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    wb.save(path)
    return len(srows)


# ---------------------------------------------------------------- the data dictionary and the file register
DICT = [
    ("glra", "TEST_TAKER_ID", "Calderwood test-taker number printed on the answer booklet"),
    ("glra", "STATE_STUDENT_ID", "State student identifier from the booklet's pre-ID label"),
    ("glra", "DISTRICT_CODE", "District of the testing school, spring 2025"),
    ("glra", "SCHOOL_CODE", "Testing school"),
    ("glra", "GRADE", "Grade tested"),
    ("glra", "FORM", "Test form"),
    ("glra", "SESSION", "Regular or Make-up"),
    ("glra", "TEST_DATE", "Date of the session"),
    ("glra", "SCORE_STATUS", "SCR scored, ABS absent, EXM exempt, INV invalidated"),
    ("glra", "RAW_SCORE", "Items correct, 0 to 30"),
    ("glra", "SCALE_SCORE", "Scale score, 150 to 450"),
    ("glra", "PERF_LEVEL", "Performance level of the scale score"),
    ("crosswalk", "TEST_TAKER_ID", "Calderwood test-taker number"),
    ("crosswalk", "STATE_STUDENT_ID", "State student identifier"),
    ("crosswalk", "MATCH_METHOD", "How the number was matched to the state identifier"),
    ("crosswalk", "MATCH_DATE", "Date of the match"),
    ("slm", "state_student_id", "State student identifier"),
    ("slm", "district_code", "District of the testing school, spring 2026"),
    ("slm", "school_code", "Testing school"),
    ("slm", "grade", "Grade tested"),
    ("slm", "session", "Main or Make-up"),
    ("slm", "test_date", "Date the session began"),
    ("slm", "status", "Scored, Absent, Exempt or Invalidated"),
    ("slm", "scale_score", "Scale score, 2100 to 2900"),
    ("slm", "performance_level", "Performance level reported with the score, 1 to 4"),
    ("bridge", "state_student_id", "State student identifier"),
    ("bridge", "district_code", "District of the school"),
    ("bridge", "school_code", "School in the study"),
    ("bridge", "order_of_tests", "Which test the school sat first"),
    ("bridge", "glra_test_date", "Date of the retiring test's session"),
    ("bridge", "glra_raw_score", "Items correct on the retiring test's grade 8 form"),
    ("bridge", "glra_scale_score", "Retiring test's scale score"),
    ("bridge", "slm_scale_score", "The student's spring 2026 grade 8 score on the Statewide Literacy Measure"),
    ("enrollment", "district_code", "District"),
    ("enrollment", "school_code", "School"),
    ("enrollment", "state_student_id", "State student identifier"),
    ("enrollment", "grade", "Grade"),
    ("enrollment", "entry_date", "First day on the school's roll"),
    ("enrollment", "exit_date", "Last day on the school's roll; blank while enrolled"),
    ("enrollment", "exit_reason", "Why the student left the school's roll"),
    ("nonpublic", "school_code", "School code assigned for state testing"),
    ("nonpublic", "school_name", "School"),
    ("nonpublic", "county", "County"),
    ("nonpublic", "located_in_district", "Public school district in which the school stands"),
    ("nonpublic", "grades_served", "Grades"),
    ("nonpublic", "affiliation", "Affiliation"),
    ("nonpublic", "state_test_participation", "Yes where the school's students sit the state tests by agreement"),
    ("finance", "district_code", "District"),
    ("finance", "district_name", "District"),
    ("finance", "fiscal_year", "State fiscal year (July to June)"),
    ("finance", "fall_2024_enrollment", "Students enrolled on October 1, 2024, all grades"),
    ("finance", "current_expenditure", "Current expenditure, dollars"),
    ("finance", "expenditure_per_pupil", "Current expenditure per enrolled student, dollars"),
    ("gradrate", "district_code", "District"),
    ("gradrate", "district_name", "District"),
    ("gradrate", "cohort", "Students who entered grade 9 in the fall named"),
    ("gradrate", "cohort_size", "Students in the cohort, adjusted for transfers in and out"),
    ("gradrate", "graduates_four_year", "Cohort students graduating within four years"),
    ("gradrate", "graduation_rate", "Four-year graduation rate, per cent"),
    ("directory", "(sheet Districts)", "District code, name, county, locale and status"),
    ("directory", "(sheet Schools)", "Public middle and high schools: code, name, district and grades"),
]


def write_dictionary(path):
    rows = [[W.FILES[k], f, d] for k, f, d in DICT]
    return P.write_csv(path, ["file_name", "field_name", "field_description"], rows)


REGISTER = [
    ("glra", W.OLD_VENDOR, "Grade-Level Reading Assessment, spring 2025, student results, grades 6 to 8", "2025-06-20"),
    ("crosswalk", W.OLD_VENDOR, "Test-taker numbers matched to state student identifiers, spring 2025", "2025-07-15"),
    ("glra_guide", W.OLD_VENDOR, "Grade-Level Reading Assessment, spring 2025, interpretive guide", "2025-08-15"),
    ("slm", W.NEW_VENDOR, "Statewide Literacy Measure, spring 2026, student results, grades 6 to 8", "2026-06-30"),
    ("slm_report", W.NEW_VENDOR, "Statewide Literacy Measure, spring 2026, grade 8 technical summary", "2026-06-26"),
    ("bridge", W.OFFICE, "Bridge study, spring 2026: paired grade 8 scores on both tests", "2026-08-05"),
    ("bridge_memo", W.OFFICE, "Bridge study memorandum", "2026-08-21"),
    ("resolution", W.BOARD, "Resolution 2026-22: cut scores for grades 3 to 7", "2026-08-13"),
    ("enrollment", "Office of Data Management", "October 2026 enrolment collection, grades 9 to 12, enrolment spells",
     "2026-11-02"),
    ("directory", "Office of Data Management", "District and school directory, 2026-27", "2026-10-01"),
    ("nonpublic", "Office of Nonpublic Schools", "Registered nonpublic schools whose students sit the state tests",
     "2026-09-15"),
    ("rules", W.LITERACY_OFFICE, "Adolescent Literacy Partnership: 2027 grant rules and applicants", "2026-10-16"),
    ("finance", "Office of School Finance", "District current expenditure per pupil, fiscal year 2025", "2026-02-27"),
    ("gradrate", "Office of Accountability", "Four-year graduation rates, class of 2025", "2026-01-30"),
    ("dict", W.LITERACY_OFFICE, "Data dictionary for this folder", "2026-11-20"),
]


def write_register(path, counts):
    rows = []
    for k, src, desc, asof in REGISTER:
        rows.append([W.FILES[k], src, desc, asof, counts.get(k, "")])
    return P.write_csv(path, ["file_name", "supplied_by", "what_it_holds", "as_of_date", "record_count"], rows)
