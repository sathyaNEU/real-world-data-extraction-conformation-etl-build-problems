"""Read the shipped pack as a solver would: the count-day roll, each student's latest grade 8 result from the two
vendors' files (the retiring test's rows keyed through the crosswalk), the bridge study's pairs, the directory's
schools and the grade 7 rows of the new test. Everything golden.py and checks.py compute comes from here."""
import csv
import os

import numpy as np

import world as W


def rows(target, key):
    with open(os.path.join(target, W.FILES[key]), encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def directory(target):
    """{school code: district code} for the directory's public schools, and {code: name} for its districts."""
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(target, W.FILES["directory"]), read_only=True)
    schools, districts = {}, {}
    for r in wb["Schools"].iter_rows(min_row=5, values_only=True):
        if r[0]:
            schools[str(r[0])] = str(r[2])
    for r in wb["Districts"].iter_rows(min_row=5, values_only=True):
        if r[0]:
            districts[str(r[0])] = str(r[1])
    return schools, districts


def roll(target, districts=None, grades=(9, 10)):
    """{state id: (district, school, grade)} for the spells covering the count day."""
    out = {}
    for r in rows(target, "enrollment"):
        g = int(r["grade"])
        if g not in grades or (districts is not None and r["district_code"] not in districts):
            continue
        if r["entry_date"] <= W.COUNT_DAY and (r["exit_date"] == "" or r["exit_date"] >= W.COUNT_DAY):
            out[r["state_student_id"]] = (r["district_code"], r["school_code"], g)
    return out


def grade8_results(target):
    """{state id: [(test, year, score), ...]} for every grade 8 session that produced a score."""
    xw = {r["TEST_TAKER_ID"]: r["STATE_STUDENT_ID"] for r in rows(target, "crosswalk")}
    out = {}
    for r in rows(target, "glra"):
        if r["GRADE"] == "8" and r["SCORE_STATUS"] == "SCR":
            sid = r["STATE_STUDENT_ID"] or xw[r["TEST_TAKER_ID"]]
            out.setdefault(sid, []).append(("old", 2025, int(r["SCALE_SCORE"])))
    for r in rows(target, "slm"):
        if r["grade"] == "8" and r["status"] == "Scored":
            out.setdefault(r["state_student_id"], []).append(("new", 2026, int(r["scale_score"])))
    return out


def population(target, districts=None):
    """(district, grade, test, year, score, state id, school) for each student on the count-day roll in grade 9 or
    10 with a grade 8 state reading result, the latest one."""
    res = grade8_results(target)
    out = []
    for sid, (d, school, g) in sorted(roll(target, districts).items()):
        if sid in res:
            test, year, sc = max(res[sid], key=lambda t: t[1])
            out.append((d, g, test, year, sc, sid, school))
    return out


def bridge(target):
    b = rows(target, "bridge")
    return (np.array([int(r["glra_scale_score"]) for r in b]), np.array([int(r["slm_scale_score"]) for r in b]),
            np.array([r["school_code"] for r in b]), np.array([r["order_of_tests"] for r in b]))


def grade7(target, by="directory", cuts="adopted"):
    """Each applicant's spring 2026 grade 7 students with a score on the new test: [at Level 1, with a score]. by:
    the directory's schools (each district's own schools) or the file's district code; cuts: the Board's adopted
    grade 7 cut or the file's performance level."""
    schools, _ = directory(target)
    merged = {sub: new for new, parts in W.MERGER.items() for sub, _, _ in parts}
    out = {d: [0, 0] for d in W.APPLICANTS}
    for r in rows(target, "slm"):
        if r["grade"] != "7" or r["status"] != "Scored":
            continue
        d = schools.get(r["school_code"]) if by == "directory" else merged.get(r["district_code"], r["district_code"])
        if d not in out:
            continue
        sc = int(r["scale_score"])
        lvl1 = sc < W.NEW_CUTS[7][0] if cuts == "adopted" else r["performance_level"] == "1"
        out[d][0] += int(lvl1)
        out[d][1] += 1
    return out
