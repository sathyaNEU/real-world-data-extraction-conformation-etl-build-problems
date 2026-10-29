"""task464 architecture 10: every graded figure, computed from the cut pack under target/.

The call: each project's 2025/26 saving, the I2R losses its replaced sections would not have incurred on the new
conductor, at each section's own half-hourly current, the net of every connection's metered flow beyond it. The asks:
each feeder's 2025/26 CI and CML under the network's reporting definitions. figures.json carries every graded figure;
ceiling.py and checks.py import the loaders and the routes from here.
"""
import csv
import datetime as dt
import json
import os
import re
import numpy as np
import openpyxl
import pyarrow.parquet as pq
import world as W

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.dirname(HERE)
TGT = os.path.join(BASE, "target")
NHH = W.NHH
UTC0 = W.UTC0
CODE_OF = {v: k for k, v in W.MODEL_CODE.items()}
FEEDER_ORDER = [f["name"] for f in W.FEEDERS]


def uk_offset(t):
    def last_sunday(y, m):
        d = dt.date(y, m, 31)
        while d.weekday() != 6:
            d -= dt.timedelta(days=1)
        return dt.datetime(d.year, d.month, d.day, 1, 0)
    return 1 if last_sunday(t.year, 3) <= t < last_sunday(t.year, 10) else 0


UTC = [UTC0 + dt.timedelta(minutes=30 * i) for i in range(NHH)]
LOCAL = [t + dt.timedelta(hours=uk_offset(t)) for t in UTC]
QUARTER = np.array([{4: 0, 5: 0, 6: 0, 7: 1, 8: 1, 9: 1, 10: 2, 11: 2, 12: 2, 1: 3, 2: 3, 3: 3}[l.month] for l in LOCAL])


# ---------------------------------------------------------------- loaders
def load_network(tgt=TGT):
    wb = openpyxl.load_workbook(os.path.join(tgt, "hv_network_model.xlsx"), read_only=True)
    secs = []
    for row in list(wb["Sections"].iter_rows(values_only=True))[1:]:
        sid, circ, frm, to, metres, cond, cons, inst, state = row
        secs.append(dict(id=sid, circuit=circ, frm=frm, to=to, metres=float(metres), code=CODE_OF[cond], state=state))
    nodes = {}
    for row in list(wb["Nodes"].iter_rows(values_only=True))[1:]:
        nid, circ, kind, ref, desc = row
        nodes[nid] = dict(circuit=circ, kind=kind, ref=ref, desc=desc)
    return secs, nodes


def asset_of(ref):
    """('S', site_no) for a substation reference written either way, ('H', agreement) for an HV customer."""
    if ref is None or str(ref).strip() == "":
        return None
    s = str(ref).strip()
    if s.startswith("CA-"):
        return ("H", s)
    m = re.fullmatch(r"(?:SS\s*)?(\d{6})", s)
    return ("S", int(m.group(1)))


def load_monitors(tgt=TGT):
    """site_no -> (P kW, Q kVAr) arrays over the 17,520 UTC half-hours; NaN where the monitor holds no record."""
    t = pq.read_table(os.path.join(tgt, "lv_monitoring_2025-26.parquet")).to_pandas()
    idx = ((t["period_start_utc"].dt.tz_convert(None) - np.datetime64(UTC0)) / np.timedelta64(30, "m")).astype(int).to_numpy()
    kw, kvar, kva, pf = (t[c].to_numpy(dtype=float) for c in ("kw", "kvar", "kva", "pf"))
    va = np.isnan(kw)
    P = np.where(va, kva * pf, kw)
    Q = np.where(va, kva * np.sqrt(np.clip(1 - pf ** 2, 0, None)), kvar)
    out = {}
    site = t["site_no"].to_numpy()
    for s in np.unique(site):
        m = site == s
        p, q = np.full(NHH, np.nan), np.full(NHH, np.nan)
        p[idx[m]] = P[m]
        q[idx[m]] = Q[m]
        out[int(s)] = (p, q)
    return out


def settlement_index(date_s, period):
    d = dt.datetime.strptime(date_s, "%Y%m%d")
    midnight_utc = d - dt.timedelta(hours=uk_offset(d - dt.timedelta(hours=1)))
    t = midnight_utc + dt.timedelta(minutes=30 * (int(period) - 1))
    return int((t - UTC0).total_seconds() // 1800)


def load_settlement(tgt=TGT):
    """mpan -> (kW, kVAr) average over each UTC half-hour."""
    rows = {}
    cache = {}
    with open(os.path.join(tgt, "hv_settlement_2025-26.csv")) as fh:
        rd = csv.reader(fh)
        next(rd)
        for mpan, d, sp, a, rq in rd:
            k = (d, sp)
            if k not in cache:
                cache[k] = settlement_index(d, sp)
            i = cache[k]
            if mpan not in rows:
                rows[mpan] = (np.zeros(NHH), np.zeros(NHH))
            rows[mpan][0][i] = 2 * float(a)
            rows[mpan][1][i] = 2 * float(rq)
    return rows


def load_registers(tgt=TGT):
    subs = {}
    with open(os.path.join(tgt, "secondary_substations.csv")) as fh:
        for r in csv.DictReader(fh):
            subs[int(r["site_no"])] = r
    hv = {}
    with open(os.path.join(tgt, "hv_connections.csv")) as fh:
        for r in csv.DictReader(fh):
            hv[r["agreement_ref"]] = r
    return subs, hv


def load(tgt=TGT):
    secs, nodes = load_network(tgt)
    subs, hv = load_registers(tgt)
    return dict(secs=secs, nodes=nodes, subs=subs, hv=hv, mon=load_monitors(tgt), settle=load_settlement(tgt))


# ---------------------------------------------------------------- the network's topology and each section's flow
def topology(pk):
    closed = [s for s in pk["secs"] if s["state"] == "Closed"]
    child = {}
    for s in closed:
        child.setdefault(s["frm"], []).append(s)
    beyond = {}

    def walk(n):
        out = [n]
        for s in child.get(n, []):
            out += walk(s["to"])
        return out
    for s in closed:
        beyond[s["id"]] = walk(s["to"])
    return closed, beyond


def connection_flows(pk, mode="net"):
    """asset -> (P, Q) half-hourly at the connection.

    mode net: monitors as recorded (a monitor with no record carries nothing) and each HV customer's import less its
    export; imports: HV customers' imports only; monitors: secondary substations only; absadd: export added to import."""
    out = {}
    for site, (p, q) in pk["mon"].items():
        out[("S", site)] = (np.nan_to_num(p), np.nan_to_num(q))
    if mode == "monitors":
        return out
    for ag, r in pk["hv"].items():
        if r["primary"] != "COPSEWELL CROSS":
            continue
        pi, qi = pk["settle"][r["import_mpan"]]
        pe, qe = pk["settle"][r["export_mpan"]] if r["export_mpan"] else (np.zeros(NHH), np.zeros(NHH))
        if mode == "net":
            out[("H", ag)] = (pi - pe, qi - qe)
        elif mode == "imports":
            out[("H", ag)] = (pi, np.clip(qi, 0, None))
        elif mode == "absadd":
            out[("H", ag)] = (pi + pe, np.abs(qi) + np.abs(qe))
    return out


_PROJ = None


def projects():
    """feeder name -> (breaker, [section ids], new conductor code, route km, existing conductor), as the brief prints."""
    global _PROJ
    if _PROJ is None:
        import sim
        feeders, nodes, sections, conns, neigh = sim.build_network()
        _PROJ = {}
        for f in feeders:
            secs = [s for s in sections if s.feeder == f["code"] and s.project]
            _PROJ[f["name"]] = (f["breaker"], [s.id for s in secs], f["new"], sum(s.metres for s in secs) / 1000.0,
                                "; ".join(sorted(set(W.CONDUCTORS[s.code][0] for s in secs))))
    return _PROJ


def section_saving_kwh(P, Q, sec, new):
    dR = (W.CONDUCTORS[sec["code"]][2] - W.CONDUCTORS[new][2]) * sec["metres"] / 1000.0
    return (P ** 2 + Q ** 2) * dR / (W.V_NOM_KV ** 2) / 1000.0 * 0.5


def savings(pk, mode="net", proj=None, reduce=None):
    """feeder -> half-hourly kWh saved, every section's own flow the sum of the connections beyond it."""
    proj = proj or projects()
    closed, beyond = topology(pk)
    byid = {s["id"]: s for s in pk["secs"]}
    fl = connection_flows(pk, mode)
    out = {}
    for name, (brk, ids, new, km, old) in proj.items():
        tot = np.zeros(NHH)
        for sid in ids:
            P, Q = np.zeros(NHH), np.zeros(NHH)
            for n in beyond[sid]:
                a = asset_of(pk["nodes"][n]["ref"])
                if a is not None and a in fl:
                    P = P + fl[a][0]
                    Q = Q + fl[a][1]
            if reduce is not None:
                P, Q = reduce(P), reduce(Q)
            tot += section_saving_kwh(P, Q, byid[sid], new)
        out[name] = tot
    return out


def summarise(hh):
    """feeder -> dict(q=[MWh x4], year MWh) from half-hourly kWh."""
    return {k: dict(q=[float(v[QUARTER == i].sum() / 1000) for i in range(4)], year=float(v.sum() / 1000)) for k, v in hh.items()}


# ---------------------------------------------------------------- the routes a solver might take (ceiling.py, checks.py)
def load_scada(tgt=TGT):
    """circuit -> (mean phase current A, MW, MVAr) arrays in the file's row order, which is the UTC order."""
    out = {}
    with open(os.path.join(tgt, "copsewell_cross_scada_2025-26.csv")) as fh:
        rd = csv.reader(fh)
        next(rd)
        for ts, circ, a, b, c, mw, mvar in rd:
            out.setdefault(circ, []).append(((float(a) + float(b) + float(c)) / 3.0, float(mw), float(mvar)))
    return {k: tuple(np.array(x) for x in zip(*v)) for k, v in out.items()}


def routes(pk, tgt=TGT):
    """feeder -> year MWh under each construction of a replaced section's current."""
    proj = projects()
    byid = {s["id"]: s for s in pk["secs"]}
    closed, beyond = topology(pk)
    sc = load_scada(tgt)
    circ_of = {}
    for name, (brk, ids, new, km, old) in proj.items():
        circ_of[name] = "%s %s" % (brk, name.upper())
    kva = {("S", int(k)): float(v["transformer_kva"]) for k, v in pk["subs"].items()}
    hv_imp = {("H", k): float(v["import_capacity_kva"] or 0) for k, v in pk["hv"].items()}
    hv_exp = {("H", k): float(v["export_capacity_kw"] or 0) for k, v in pk["hv"].items()}
    out = {}

    def share_route(cap):
        res = {}
        for name, (brk, ids, new, km, old) in proj.items():
            I = sc[circ_of[name]][0]
            top = [s for s in closed if s["circuit"] == circ_of[name] and pk["nodes"][s["frm"]]["kind"] == "Busbar"][0]
            def capof(sid):
                return sum(cap.get(asset_of(pk["nodes"][n]["ref"]), 0.0) for n in beyond[sid])
            tot = capof(top["id"])
            v = 0.0
            for sid in ids:
                sec = byid[sid]
                dR = (W.CONDUCTORS[sec["code"]][2] - W.CONDUCTORS[new][2]) * sec["metres"] / 1000.0
                v += (3 * (I * capof(sid) / tot) ** 2 * dR * 0.5 / 1000.0).sum() / 1000.0
            res[name] = float(v)
        return res
    head = {}
    head_s = {}
    for name, (brk, ids, new, km, old) in proj.items():
        I, mw, mvar = sc[circ_of[name]]
        Is = np.sqrt(mw ** 2 + mvar ** 2) * 1000 / (np.sqrt(3) * W.V_NOM_KV)
        a = b = 0.0
        for sid in ids:
            sec = byid[sid]
            dR = (W.CONDUCTORS[sec["code"]][2] - W.CONDUCTORS[new][2]) * sec["metres"] / 1000.0
            a += (3 * I ** 2 * dR * 0.5 / 1000.0).sum() / 1000.0
            b += (3 * Is ** 2 * dR * 0.5 / 1000.0).sum() / 1000.0
        head[name], head_s[name] = float(a), float(b)
    out["head_current"] = head
    out["head_mw_mvar"] = head_s
    out["kva_share"] = share_route(kva)
    out["kva_share_hv_import"] = share_route({**kva, **hv_imp})
    out["kva_share_hv_all"] = share_route({**kva, **{k: hv_imp[k] + hv_exp[k] for k in hv_imp}})
    for mode, lab in (("monitors", "monitors_only"), ("imports", "loads_only"), ("absadd", "export_added"), ("net", "net")):
        out[lab] = {k: float(v.sum() / 1000) for k, v in savings(pk, mode, proj).items()}
    out["net_active_only"] = {k: float(v.sum() / 1000) for k, v in
                              savings(pk, "net", proj, reduce=None).items()}
    pk2 = dict(pk)
    pk2["mon"] = {s: (p, np.zeros_like(q)) for s, (p, q) in pk["mon"].items()}
    pk2["settle"] = {m: (p, np.zeros_like(q)) for m, (p, q) in pk["settle"].items()}
    out["net_active_only"] = {k: float(v.sum() / 1000) for k, v in savings(pk2, "net", proj).items()}
    day = np.repeat(np.arange(NHH // 48), 48)

    def daily(x):
        return np.repeat(np.bincount(day, weights=x) / 48.0, 48)
    out["net_daily_mean"] = {k: float(v.sum() / 1000) for k, v in savings(pk, "net", proj, reduce=daily).items()}
    return out


def winner(d):
    o = sorted(d, key=lambda k: -d[k])
    return o[0], o[1], d[o[0]] / d[o[1]]


# ---------------------------------------------------------------- the interruption record
def load_interruptions(tgt=TGT):
    incs = {}
    with open(os.path.join(tgt, "interruptions_2023-2026.csv")) as fh:
        for r in csv.DictReader(fh):
            i = incs.setdefault(r["reference"], dict(ref=r["reference"], circuit=r["circuit"], voltage=r["voltage"],
                                                     category=r["category"], cause=r["cause"],
                                                     start=dt.datetime.strptime(r["incident_start"], "%Y-%m-%d %H:%M"),
                                                     steps=[]))
            fmt = "%Y-%m-%d %H:%M:%S" if len(r["step_time"]) > 16 else "%Y-%m-%d %H:%M"
            t = dt.datetime.strptime(r["step_time"], fmt)
            n = int(r["customers_off"]) if r["customers_off"] else -int(r["customers_restored"])
            i["steps"].append((t, int(r["site_no"]), n))
    cust = {}
    with open(os.path.join(tgt, "customers_by_substation_2023-2026.csv")) as fh:
        for r in csv.DictReader(fh):
            cust.setdefault(int(r["site_no"]), {})[r["month_end"]] = (int(r["customers"]), r["circuit"].replace(" (HV)", ""))
    wb = openpyxl.load_workbook(os.path.join(tgt, "generator_hire_2025-26.xlsx"), read_only=True)
    hires = []
    for row in list(wb["Hire log"].iter_rows(values_only=True))[1:]:
        hires.append(dict(site=int(row[4]), code=row[6], connected=dt.datetime.strptime(row[8], "%d/%m/%Y %H:%M"),
                          disconnected=dt.datetime.strptime(row[9], "%d/%m/%Y %H:%M")))
    return list(incs.values()), cust, hires


STORM = (dt.datetime(2026, 1, 6, 15, 0), dt.datetime(2026, 1, 7, 23, 59))


def interruption_measures(incs, cust, hires, storm_circuits, *, generators=True, dedupe=True, short_rule=True,
                          prearranged_half=True, exceptional_rule=True, transmission_rule=True, denominator="2026-03-31"):
    lo, hi = dt.datetime(2025, 4, 1), dt.datetime(2026, 4, 1)
    site_circ = {s: v[max(v)][1] for s, v in cust.items()}
    # each site's outage spans in the HV incidents (who was already off)
    hv_off = {}
    for inc in incs:
        if inc["voltage"] == "LV":
            continue
        per = {}
        for t, s, n in inc["steps"]:
            per.setdefault(s, []).append((t, n))
        for s, ev in per.items():
            ev.sort()
            on = [t for t, n in ev if n < 0][0]
            hv_off.setdefault(s, []).append((ev[0][0], on))
    ci, cml = {}, {}
    for inc in incs:
        if not (lo <= inc["start"] < hi):
            continue
        if transmission_rule and inc["voltage"] == "EHV":
            continue
        if exceptional_rule and STORM[0] <= inc["start"] <= STORM[1] and inc["circuit"] in storm_circuits:
            continue
        if dedupe and inc["voltage"] == "LV":
            t0 = min(t for t, s, n in inc["steps"])
            s0 = inc["steps"][0][1]
            if any(a <= t0 < b for a, b in hv_off.get(s0, [])):
                continue
        w = 0.5 if (prearranged_half and inc["category"] == "Pre-arranged") else 1.0
        per = {}
        for t, s, n in inc["steps"]:
            per.setdefault(s, []).append((t, n))
        for s, ev in per.items():
            ev.sort()
            off_t, n = ev[0]
            on_t = [t for t, x in ev if x < 0][0]
            if generators:
                held = False
                for h in hires:
                    if h["site"] != s:
                        continue
                    if h["connected"] <= off_t and h["disconnected"] >= on_t:
                        held = True                       # kept on supply throughout: never lost it
                    elif off_t <= h["connected"] < on_t:
                        on_t = h["connected"]             # supply restored when the generator was connected
                if held:
                    continue
            mins = (on_t - off_t).total_seconds() / 60.0
            if short_rule and mins < 3:
                continue
            c = site_circ[s]
            ci[c] = ci.get(c, 0.0) + w * n
            cml[c] = cml.get(c, 0.0) + w * n * mins
    den = {}
    for s, v in cust.items():
        if denominator in v:
            c = v[denominator][1]
            den[c] = den.get(c, 0) + v[denominator][0]
    return {c: dict(ci=100.0 * ci.get(c, 0.0) / den[c], cml=cml.get(c, 0.0) / den[c], customers=den[c]) for c in FEEDER_ORDER}


def storm_circuits_from_pdf(tgt=TGT):
    """The circuits the accepted exceptional event names, read from the definitions PDF."""
    from pypdf import PdfReader
    txt = " ".join(p.extract_text() for p in PdfReader(os.path.join(tgt, "iis_reporting_definitions_2025-26.pdf")).pages)
    txt = re.sub(r"\s+", " ", txt)
    return [n for n in FEEDER_ORDER if n in txt.split("Exceptional event accepted for 2025/26", 1)[1]]


# ---------------------------------------------------------------- the figures
def figures(write=True, tgt=TGT):
    pk = load(tgt)
    proj = projects()
    s = summarise(savings(pk, "net", proj))
    order = sorted(s, key=lambda k: -s[k]["year"])
    call, runner = order[0], order[1]
    incs, cust, hires = load_interruptions(tgt)
    sc = storm_circuits_from_pdf(tgt)
    iis = interruption_measures(incs, cust, hires, sc)
    rt = routes(pk, tgt)
    rungs = {}
    for lab, d in rt.items():
        w0, w1, m = winner(d)
        rungs[lab] = dict(winner=w0, second=w1, margin=round(m, 3), mwh={k: round(v, 1) for k, v in d.items()},
                          kestle_rank=sorted(d, key=lambda k: -d[k]).index(call) + 1)
    fig = dict(
        rungs=rungs,
        call=call, runner_up=runner,
        gap_mwh=round(s[call]["year"], 1) - round(s[runner]["year"], 1),
        losses_value_gbp_mwh=W.LOSSES_VALUE_GBP_MWH,
        projects={k: dict(breaker=proj[k][0], sections=proj[k][1], new=proj[k][2], km=round(proj[k][3], 2), old=proj[k][4],
                          new_desc=W.CONDUCTORS[proj[k][2]][0],
                          year_mwh=round(s[k]["year"], 1),
                          value_gbp=int(round(s[k]["year"] * W.LOSSES_VALUE_GBP_MWH, -2)),
                          dR_ohm=round(sum((W.CONDUCTORS[x["code"]][2] - W.CONDUCTORS[proj[k][2]][2]) * x["metres"] / 1000.0
                                           for x in pk["secs"] if x["id"] in proj[k][1]), 3),
                          raw=dict(q=s[k]["q"], year=s[k]["year"]))
                  for k in FEEDER_ORDER},
        order=order,
        iis={k: dict(ci=round(v["ci"], 1), cml=round(v["cml"], 1), customers=v["customers"],
                     raw=dict(ci=v["ci"], cml=v["cml"])) for k, v in iis.items()},
        storm_circuits=sc,
        conventions=dict(gbp_rounding=100, ci_per_customers=100, v_kv=W.V_NOM_KV),
    )
    if write:
        json.dump(fig, open(os.path.join(HERE, "figures.json"), "w"), indent=1, sort_keys=True)
    return fig


if __name__ == "__main__":
    import sys
    f = figures(write=False, tgt=sys.argv[1] if len(sys.argv) > 1 else TGT)
    for k in f["order"]:
        p = f["projects"][k]
        print("%-14s %s  year %6.1f  GBP %d" % (k, p["q_mwh"], p["year_mwh"], p["value_gbp"]))
    print("call", f["call"], "runner", f["runner_up"], "gap", f["gap_mwh"])
    for k, v in f["iis"].items():
        print("%-14s CI %6.1f  CML %6.1f  customers %d" % (k, v["ci"], v["cml"], v["customers"]))
