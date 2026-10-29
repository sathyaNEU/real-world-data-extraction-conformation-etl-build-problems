"""task464 architecture 10: write the data files of the pack as each system exports them.

  copsewell_cross_scada_2025-26.csv          primary substation SCADA, feeder breakers, UK local time, half-hour ending
  lv_monitoring_2025-26.parquet            secondary substation monitors (Copsewell Cross and Feltmoor), UTC period start
  hv_settlement_2025-26.csv                half-hourly settlement data for HV-connected customers, settlement periods
  hv_network_model.xlsx                    the network model: sections and nodes, normally-open points included
  secondary_substations.csv                substation register
  hv_connections.csv                       HV connections register
  interruptions_2023-2026.csv              interruption log (control-system HV incidents and work-management LV jobs)
  customers_by_substation_2023-2026.csv    customers connected at each substation at each month end
  generator_hire_2025-26.xlsx              mobile generator hire and fuel log
"""
import csv
import datetime as dt
import io
import os
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment
import pyarrow as pa
import pyarrow.parquet as pq
import world as W
import sim as S

OUT = None


def path(name):
    return os.path.join(OUT, name)


def _r(tag):
    return S.rng("pack:" + tag)


# ---------------------------------------------------------------- the network's identifiers as each system writes them
def site_text(site, r):
    return ("SS %d" % site) if r.random() < 0.55 else str(site)


def node_text(nid, r):
    return nid if r.random() < 0.7 else nid[:2] + " " + nid[2:]


ENERGISED = {"Hob Moor Solar Park": dt.date(2023, 6, 28), "Crane Hill wind turbine": dt.date(2019, 11, 4),
             "Ashlade Community Hospital": dt.date(2012, 3, 19)}


def hv_identity(w):
    """Each HV customer's connection agreement reference and energisation date (the agreement is numbered in the year
    the connection was offered, a year or two before it was energised)."""
    r = _r("hv-identity")
    out = {}
    used = set()
    for c in w["conns"] + w["neigh"]:
        if c.kind != "H":
            continue
        en = ENERGISED.get(c.name) or dt.date(int(r.integers(1998, 2016)), int(r.integers(1, 13)), int(r.integers(1, 28)))
        yr = en.year - int(r.integers(1, 3))
        while True:
            num = int(r.integers(120, 2900))
            if num not in used:
                used.add(num)
                break
        out[c.key] = dict(agreement="CA-%d-%04d" % (yr, num), energised=en)
    return out


def write_network(w):
    r = _r("network")
    names = {f["code"]: f["name"] for f in w["feeders"]}
    brk = {f["code"]: f["breaker"] for f in w["feeders"]}
    bykey = {c.key: c for c in w["conns"]}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sections"
    hdr = ["Section ID", "Circuit", "From node", "To node", "Length (m)", "Conductor", "Construction", "Installed", "Normal state"]
    ws.append(hdr)
    yr = {}
    rows = []
    for s in w["sections"]:
        code = W.CONDUCTORS[s.code]
        old = ("PILC" in s.code) or s.code in ("OH-CU-0.04", "OH-CU-0.1", "OH-ACSR-50")
        y = int(r.integers(1958, 1979)) if old else int(r.integers(1991, 2021))
        if any(k == "S:Redburn Meadows" for k in s.down) and len(s.down) == 1:
            y = 2025                                   # the new estate's spur
        circ = "%s %s" % (brk[s.feeder], names[s.feeder].upper())
        rows.append([s.id, circ, s.frm, s.to, s.metres, W.MODEL_CODE[s.code], "Underground cable" if code[1] == "UG" else "Overhead line",
                     y, "Open" if getattr(s, "normally_open", False) else "Closed"])
    order = list(range(len(rows)))
    r.shuffle(order)
    for i in sorted(order, key=lambda i: (rows[i][1], rows[i][0])):
        ws.append(rows[i])
    secmap = {s.id: s for s in w["sections"]}
    ws2 = wb.create_sheet("Nodes")
    ws2.append(["Node ID", "Circuit", "Node type", "Asset reference", "Description"])
    for nid, n in w["nodes"].items():
        circ = "%s %s" % (brk[n.feeder], names[n.feeder].upper())
        ref, desc = "", ""
        if n.kind == "Busbar":
            desc = "Copsewell Cross 11kV board, %s" % brk[n.feeder]
        elif getattr(n, "asset", None):
            c = bykey[n.asset]
            if c.kind == "S":
                ref = site_text(c.site, r)
                desc = c.name
            else:
                ref = w["hv"][c.key]["agreement"]
                desc = c.name
        elif n.kind == "RMU":
            par = secmap.get(n.parent_section)
            oh = par is not None and W.CONDUCTORS[par.code][1] == "OH"
            if oh:
                desc = ["Pole-mounted switch", "Pole-mounted recloser", "Pole-mounted switch", "Air break switch"][int(r.integers(0, 4))]
            else:
                desc = ["Ring main unit", "Ring main unit", "Section switch", "Ring main unit (remote control)"][int(r.integers(0, 4))]
        rows2 = [nid, circ, n.kind if n.kind != "RMU" else "Switching point", ref, desc]
        ws2.append(rows2)
    for wsx in (ws, ws2):
        for c in wsx[1]:
            c.font = Font(bold=True)
        wsx.freeze_panes = "A2"
        for col, wd in zip("ABCDEFGHI", [12, 22, 11, 11, 11, 18, 18, 10, 13]):
            wsx.column_dimensions[col].width = wd
    wb.save(path("hv_network_model.xlsx"))
    return len(rows)


def write_registers(w):
    r = _r("registers")
    names = {f["code"]: f["name"] for f in w["feeders"]}
    subs = [c for c in w["conns"] + w["neigh"] if c.kind == "S"]
    monitor_ids = sorted(set(int(x) for x in r.choice(np.arange(10000, 99999), len(subs) + 10, replace=False)))
    r.shuffle(monitor_ids)
    mids = {}
    with open(path("secondary_substations.csv"), "w", newline="") as fh:
        wr = csv.writer(fh)
        wr.writerow(["site_no", "site_name", "primary", "circuit", "transformer_kva", "monitor_id", "monitor_installed", "network_node"])
        rows = []
        for i, c in enumerate(subs):
            mid = "LVM-%05d" % monitor_ids[i]
            mids[c.key] = mid
            prim = "COPSEWELL CROSS" if c.feeder != "LUD" else "FELTMOOR"
            if c.feeder == "LUD":
                j = [x[0] for x in W.NEIGHBOUR_SUBS].index(c.name)
                circ = "Feltmoor Town" if j < 5 else ("Furnace Lane" if j < 10 else "Saltbrook")
            else:
                circ = names[c.feeder]
            inst = dt.date(2024, 1, 15) + dt.timedelta(days=int(r.integers(0, 210)))
            if c.name == "Redburn Meadows":            # the new estate's monitor went in with its substation
                inst = W.REDBURN_MEADOWS_ENERGISED
            node = node_text(c.host, r) if c.host else ""
            rows.append([c.site, c.name, prim, circ, c.kva, mid, inst.strftime("%d/%m/%Y"), node])
        rows.sort(key=lambda x: x[0])
        for row in rows:
            wr.writerow(row)
    hv = [c for c in w["conns"] + w["neigh"] if c.kind == "H"]
    with open(path("hv_connections.csv"), "w", newline="") as fh:
        wr = csv.writer(fh)
        wr.writerow(["agreement_ref", "customer", "primary", "circuit", "connection_type", "technology",
                     "import_capacity_kva", "export_capacity_kw", "import_mpan", "export_mpan", "energised", "network_node"])
        rows = []
        for c in hv:
            k = c.hvkind
            tech = {"solar": "Solar PV", "wind": "Wind", "chp": "Gas CHP", "demand": ""}[k]
            ctype = {"solar": "Generation", "wind": "Generation", "chp": "Demand and generation", "demand": "Demand"}[k]
            imp_kva = {"solar": 60, "wind": 40, "chp": 1500, "demand": int(round(c.params.get("mean_kw", 400) * 1.9, -1))}[k]
            exp_kw = {"solar": c.params.get("cap_kw", 0), "wind": c.params.get("cap_kw", 0), "chp": 1000, "demand": ""}[k]
            prim = "COPSEWELL CROSS" if c.feeder != "LUD" else "FELTMOOR"
            circ = names.get(c.feeder, "Feltmoor Town")
            en = w["hv"][c.key]["energised"].isoformat()
            rows.append([w["hv"][c.key]["agreement"], c.name, prim, circ, ctype, tech, imp_kva, exp_kw, "%s" % ("40%011d" % c.mpan_import),
                         ("40%011d" % c.mpan_export) if c.mpan_export else "", en, c.host or ""])
        rows.sort(key=lambda x: x[0])
        for row in rows:
            wr.writerow(row)
    return mids


# ---------------------------------------------------------------- half-hourly records
def write_monitors(w, mids):
    r = _r("monitors")
    fl = w["flows"]
    subs = [c for c in w["conns"] + w["neigh"] if c.kind == "S"]
    ts = np.array([np.datetime64(t) for t in S.UTC], dtype="datetime64[s]")
    cols = {"monitor_id": [], "site_no": [], "period_start_utc": [], "kw": [], "kvar": [], "kva": [], "pf": []}
    recorded = {}
    for c in subs:
        p, q = fl[c.key]["p"], fl[c.key]["q"]
        keep = np.ones(S.N, bool)
        if c.name in W.MONITOR_SWAPS:
            d0, days = W.MONITOR_SWAPS[c.name]
            lo = np.datetime64(dt.datetime(d0.year, d0.month, d0.day, 9, 30))
            keep &= ~((ts >= lo) & (ts < lo + np.timedelta64(days * 86400 + 4 * 3600, "s")))
        if c.name == "Redburn Meadows":
            keep &= S.LOCAL_DATE >= W.REDBURN_MEADOWS_ENERGISED
        if c.key in w["onf"]:
            keep &= w["onf"][c.key] > 0          # a monitor off supply for the whole half-hour records nothing
        kw = np.round(p, 1)
        kvar = np.round(q, 1)
        vendor_a = (c.name in W.VENDOR_A_SUBS) & (S.LOCAL_DATE < W.VENDOR_A_UNTIL)
        kva = np.round(np.sqrt(p ** 2 + q ** 2), 1)
        pf = np.round(np.where(kva > 0, np.abs(p) / np.maximum(kva, 1e-9), 1.0), 3)
        recorded[c.key] = dict(keep=keep, vendor_a=vendor_a, kw=kw, kvar=kvar, kva=kva, pf=pf)
        idx = np.nonzero(keep)[0]
        cols["monitor_id"].extend([mids[c.key]] * len(idx))
        cols["site_no"].extend([c.site] * len(idx))
        cols["period_start_utc"].append(ts[idx])
        va = vendor_a[idx]
        cols["kw"].append(np.where(va, np.nan, kw[idx]))
        cols["kvar"].append(np.where(va, np.nan, kvar[idx]))
        cols["kva"].append(np.where(va, kva[idx], np.nan))
        cols["pf"].append(np.where(va, pf[idx], np.nan))
    tbl = pa.table({
        "monitor_id": pa.array(cols["monitor_id"], pa.string()),
        "site_no": pa.array(cols["site_no"], pa.int32()),
        "period_start_utc": pa.array(np.concatenate(cols["period_start_utc"]), pa.timestamp("s", tz="UTC")),
        "kw": pa.array(np.concatenate(cols["kw"]), pa.float32()),
        "kvar": pa.array(np.concatenate(cols["kvar"]), pa.float32()),
        "kva": pa.array(np.concatenate(cols["kva"]), pa.float32()),
        "pf": pa.array(np.concatenate(cols["pf"]), pa.float32()),
    })
    pq.write_table(tbl, path("lv_monitoring_2025-26.parquet"), compression="zstd", row_group_size=250000,
                   use_dictionary=["monitor_id"], write_statistics=False)
    return tbl.num_rows, recorded


def settlement_periods():
    """Each UTC half-hour's settlement date and period (1 to 48, 46 or 50 on the clock-change days)."""
    sd, sp = [], []
    cur, k = None, 0
    for l in S.LOCAL:
        d = l.date()
        if d != cur:
            cur, k = d, 0
        k += 1
        sd.append(d)
        sp.append(k)
    return sd, sp


def write_settlement(w):
    fl = w["flows"]
    sd, sp = settlement_periods()
    hv = [c for c in w["conns"] + w["neigh"] if c.kind == "H"]
    rows = []
    for c in hv:
        d = fl[c.key]
        imp = d["imp"] if d["imp"] is not None else np.clip(d["p"], 0, None)
        exp = d["exp"] if d.get("exp") is not None else np.zeros(S.N)
        qimp = d.get("qimp", np.clip(d["q"], 0, None))
        qexp = d.get("qexp", np.zeros(S.N))
        m_i = "40%011d" % c.mpan_import
        for i in range(S.N):
            rows.append((m_i, sd[i].strftime("%Y%m%d"), sp[i], round(float(imp[i]) * 0.5, 2), round(float(qimp[i] - qexp[i]) * 0.5, 2)))
        if c.mpan_export:
            m_e = "40%011d" % c.mpan_export
            for i in range(S.N):
                rows.append((m_e, sd[i].strftime("%Y%m%d"), sp[i], round(float(exp[i]) * 0.5, 2), 0.0))
    rows.sort(key=lambda x: (x[1], x[2], x[0]))
    with open(path("hv_settlement_2025-26.csv"), "w", newline="") as fh:
        wr = csv.writer(fh)
        wr.writerow(["mpan", "settlement_date", "settlement_period", "active_kwh", "reactive_kvarh"])
        for row in rows:
            wr.writerow(["%s" % row[0], row[1], row[2], "%.2f" % row[3], "%.2f" % row[4]])
    return len(rows)


def write_scada(w):
    """Feeder breaker metering: amps per phase, MW and MVAr, averaged over the half-hour, stamped at its end in local time."""
    r = _r("scada")
    fl = w["flows"]
    out = []
    for f in w["feeders"]:
        top = [s for s in w["sections"] if s.feeder == f["code"] and s.frm == f["bus"]][0]
        P = sum(fl[k]["p"] for k in top.down)
        Q = sum(fl[k]["q"] for k in top.down)
        loss_kw = np.zeros(S.N)
        for s in w["sections"]:
            if s.feeder != f["code"] or getattr(s, "normally_open", False):
                continue
            Ps = sum(fl[k]["p"] for k in s.down)
            Qs = sum(fl[k]["q"] for k in s.down)
            loss_kw += (Ps ** 2 + Qs ** 2) * S.ohms(s) / (W.V_NOM_KV ** 2) / 1000.0
        tx = 0.0
        for k in top.down:
            c = [c for c in w["conns"] if c.key == k][0]
            if c.kind == "S":
                tx = tx + 0.0016 * c.kva + 0.011 * (fl[k]["p"] ** 2 + fl[k]["q"] ** 2) / c.kva
        Pm = (P + loss_kw + tx) * (1 + r.normal(0, 0.009)) * (1 + r.normal(0, 0.004, S.N))
        Qm = (Q + 0.6 * loss_kw) * (1 + r.normal(0, 0.012)) * (1 + r.normal(0, 0.006, S.N))
        Sm = np.sqrt(Pm ** 2 + Qm ** 2)
        Iavg = Sm / (np.sqrt(3) * W.V_NOM_KV)
        imb = r.normal(0, 0.025, 3)
        imb -= imb.mean()
        circ = "%s %s" % (f["breaker"], f["name"].upper())
        for i in range(S.N):
            end_local = S.LOCAL[i] + dt.timedelta(minutes=30)
            out.append((end_local, S.UTC[i], circ, Iavg[i] * (1 + imb[0]), Iavg[i] * (1 + imb[1]), Iavg[i] * (1 + imb[2]),
                        Pm[i] / 1000, Qm[i] / 1000))
    out.sort(key=lambda x: (x[1], x[2]))
    with open(path("copsewell_cross_scada_2025-26.csv"), "w", newline="") as fh:
        wr = csv.writer(fh)
        wr.writerow(["timestamp", "circuit", "ia_a", "ib_a", "ic_a", "mw", "mvar"])
        for row in out:
            wr.writerow([row[0].strftime("%d/%m/%Y %H:%M"), row[2], "%.1f" % row[3], "%.1f" % row[4], "%.1f" % row[5],
                         "%.3f" % row[6], "%.3f" % row[7]])
    return len(out)


# ---------------------------------------------------------------- the interruption record
def write_reliability(w, rel):
    r = _r("reliability")
    conns = w["conns"]
    site = {c.key: c.site for c in conns}
    names = {f["code"]: f["name"] for f in w["feeders"]}
    # customers at month end
    months = sorted(next(iter(rel["counts"].values())).keys())
    with open(path("customers_by_substation_2023-2026.csv"), "w", newline="") as fh:
        wr = csv.writer(fh)
        wr.writerow(["month_end", "site_no", "circuit", "connection", "customers"])
        rows = []
        for c in conns:
            for m in months:
                me = dt.date(m.year + (m.month == 12), m.month % 12 + 1, 1) - dt.timedelta(days=1)
                n = rel["counts"][c.key][m]
                if c.name == "Redburn Meadows" and n == 0:
                    continue
                rows.append((me, c.site, names[c.feeder], "LV" if c.kind == "S" else "HV", n))
        rows.sort(key=lambda x: (x[0], x[1]))
        for row in rows:
            wr.writerow([row[0].isoformat(), row[1], row[2], row[3], row[4]])
    # interruption log: one row per site per step, in UK local time to the minute (seconds on the transient records)
    with open(path("interruptions_2023-2026.csv"), "w", newline="") as fh:
        wr = csv.writer(fh)
        wr.writerow(["reference", "circuit", "voltage", "category", "cause", "incident_start", "step", "step_time",
                     "site_no", "customers_off", "customers_restored"])
        nrow = 0
        for inc in rel["incidents"]:
            steps = sorted(inc.steps, key=lambda s: (s[0], -s[2], site.get(s[1], 0)))
            times = sorted(set(s[0] for s in steps))
            for (tm, kk, ch) in steps:
                stepno = times.index(tm)
                fmt = "%Y-%m-%d %H:%M:%S" if inc.cause.startswith("Transient") else "%Y-%m-%d %H:%M"
                wr.writerow([inc.ref, inc.circuit, inc.voltage, inc.category, inc.cause, inc.start.strftime("%Y-%m-%d %H:%M"),
                             stepno, tm.strftime(fmt), site[kk], ch if ch > 0 else "", -ch if ch < 0 else ""])
                nrow += 1
    # generator hire and fuel log: a cost record kept by the network's generation team
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hire log"
    ws.append(["Hire order", "Unit", "Rating (kVA)", "Site", "Site no", "Deployment job", "Cost code", "On hire",
               "Connected", "Disconnected", "Off hire", "Fuel (litres)", "Charge (GBP)"])
    subs = {c.key: c for c in conns}
    classes = [200, 315, 500, 800, 1000]
    units = {}
    for i, u in enumerate(sorted(rel["gen_units"])):
        units[u] = classes[[0, 1, 2, 2, 3, 3, 4, 1, 2, 3, 4, 2, 3, 4][i % 14]]
    busy = {u: [] for u in units}
    gens = [g for g in rel["gens"] if dt.datetime(2025, 4, 1) <= g["connected"] < dt.datetime(2026, 4, 1)]
    gens.sort(key=lambda g: (g["connected"], subs[g["site"]].site))
    ho = sorted(set(int(x) for x in r.choice(np.arange(70000, 99999), len(gens) + 5, replace=False)))
    jobs = sorted(set(int(x) for x in r.choice(np.arange(5100000, 5999999), len(gens) + 5, replace=False)))
    rate = {200: 165, 315: 220, 500: 305, 800: 410, 1000: 495}
    for i, g in enumerate(gens):
        c = subs[g["site"]]
        need = 1.45 * c.mean_kw
        cls = next((k for k in classes if k >= need), 1000)
        onh = g["connected"] - dt.timedelta(minutes=int(r.integers(40, 150)))
        offh = g["disconnected"] + dt.timedelta(minutes=int(r.integers(30, 240)))
        free = [u for u in sorted(units) if units[u] >= cls and all(offh <= a or onh >= b for a, b in busy[u])]
        free.sort(key=lambda u: (units[u], u))
        unit = free[int(r.integers(0, min(2, len(free))))]
        busy[unit].append((onh, offh))
        kva = units[unit]
        hours = (g["disconnected"] - g["connected"]).total_seconds() / 3600
        fuel = int(round(hours * kva * 0.19 * (0.8 + 0.35 * r.random())))
        days = max(1, (offh.date() - onh.date()).days + 1)
        charge = round(days * rate[kva] + fuel * 1.21, 2)
        ws.append(["GH%d" % ho[i], unit, kva, c.name, c.site, "WO%d" % jobs[i], "FLT-HV" if g["kind"] == "fault" else "PLN-HV",
                   onh.strftime("%d/%m/%Y %H:%M"), g["connected"].strftime("%d/%m/%Y %H:%M"),
                   g["disconnected"].strftime("%d/%m/%Y %H:%M"), offh.strftime("%d/%m/%Y %H:%M"), fuel, charge])
    for cc in ws[1]:
        cc.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for col, wd in zip("ABCDEFGHIJKLM", [11, 8, 11, 26, 9, 13, 10, 17, 17, 17, 17, 12, 12]):
        ws.column_dimensions[col].width = wd
    ws2 = wb.create_sheet("Units")
    ws2.append(["Unit", "Rating (kVA)", "Supplier", "Fuel"])
    sup = ["Marchmont Plant", "Severnside Hire", "Fenwick Generators"]
    for i, u in enumerate(sorted(units)):
        ws2.append([u, units[u], sup[int(r.integers(0, 3))], "HVO" if r.random() < 0.8 else "Diesel"])
    for cc in ws2[1]:
        cc.font = Font(bold=True)
    wb.save(path("generator_hire_2025-26.xlsx"))
    return nrow, len(gens)


def write_all(w, rel, out):
    global OUT
    OUT = out
    st = {}
    w["hv"] = hv_identity(w)
    st["sections"] = write_network(w)
    mids = write_registers(w)
    st["monitor_rows"], rec = write_monitors(w, mids)
    st["settlement_rows"] = write_settlement(w)
    st["scada_rows"] = write_scada(w)
    st["interruption_rows"], st["hires"] = write_reliability(w, rel)
    w["recorded"] = rec
    w["monitor_ids"] = mids
    return st
