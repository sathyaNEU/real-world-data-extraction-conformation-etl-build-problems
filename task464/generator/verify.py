"""task464 architecture 10: recompute every graded figure from target/ alone and check the goldens and the write-up.

    python3 generator/verify.py

Reads only the shipped bundle (the brief, the conductor data sheet and the reporting definitions are read as PDFs) and
shares no code with answers.py. Each reader is checked against a positive control before its output is used.
"""
import csv
import datetime as dt
import json
import os
import re
import sys
import numpy as np
import openpyxl
import pyarrow.parquet as pq
from pypdf import PdfReader

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
T = os.path.join(BASE, "target")
G = os.path.join(BASE, "golden")
FAIL = []


def ok(cond, msg):
    print(("PASS  " if cond else "FAIL  ") + msg)
    if not cond:
        FAIL.append(msg)


def pdf_text(name):
    return re.sub(r"[ \t]+", " ", "\n".join(p.extract_text() for p in PdfReader(os.path.join(T, name)).pages))


def utc_index(t_utc, start):
    return int((t_utc - start).total_seconds() // 1800)


def bst(t):
    """True when the UTC instant t falls in British Summer Time."""
    def last_sun(y, m):
        d = dt.date(y, m, 31)
        while d.weekday() != 6:
            d -= dt.timedelta(days=1)
        return dt.datetime(d.year, d.month, d.day, 1)
    return last_sun(t.year, 3) <= t < last_sun(t.year, 10)


def main():
    # ---------------------------------------------------------------- the documents
    cat = pdf_text("hv_conductor_data_EDS-02-0310.pdf")
    lines = [l.strip() for l in cat.split("Conductors", 1)[1].split("Notes", 1)[0].splitlines() if l.strip()]
    R20, DESC = {}, {}
    i = 0
    while i < len(lines):
        if re.match(r"^(PILC|XLPE|OHL) ", lines[i]):
            code = lines[i]
            j = i + 1
            desc = []
            while not lines[j].startswith(("Cable,", "Overhead,")):
                desc.append(lines[j])
                j += 1
            while not re.fullmatch(r"\d+\.\d+", lines[j]):
                j += 1
            R20[code] = float(lines[j])
            DESC[code] = re.sub(r"\s+", " ", " ".join(desc))
            i = j
        i += 1
    ok(len(R20) == 14 and R20.get("OHL 0.1 CU") == 0.276 and R20.get("PILC 0.3 CU") == 0.0915,
       "conductor data sheet read: %d conductors, OHL 0.1 CU 0.276 and PILC 0.3 CU 0.0915 ohm/km at 20 C" % len(R20))

    brief = pdf_text("losses_programme_2027-28_brief.pdf")
    proj_txt = brief.split("The projects", 1)[1].split("How each project is sized", 1)[0]
    flat = re.sub(r"\s+", " ", proj_txt)
    feeders = ["Copsewell Town", "Ashlade", "Oxenford Lane", "Kestle", "Pennick Hill", "Redburn", "Stowmere", "Ingle Road"]
    pos = [(flat.index(" %s CB " % f), f) for f in feeders]
    pos.sort()
    PROJ = {}
    for n, (p, f) in enumerate(pos):
        chunk = flat[p:pos[n + 1][0]] if n + 1 < len(pos) else flat[p:]
        ids = re.findall(r"HV-\d+", chunk)
        tail = chunk.split(ids[-1], 1)[1]
        new = [c for c, d in DESC.items() if tail.rstrip().endswith(d)]
        PROJ[f] = (ids, new[0] if len(new) == 1 else None)
    ok(len(PROJ) == 8 and all(v[1] for v in PROJ.values()) and len(PROJ["Kestle"][0]) == 3,
       "brief's project table read: 8 projects, each replacement conductor matched to one data sheet entry")
    m = re.search(r"losses value of GBP (\d+\.\d+) per\s*MWh", brief)
    VALUE = float(m.group(1))
    ok(VALUE == 64.80 and "11 kV nominal" in re.sub(r"\s+", " ", brief) and "20°C" in brief,
       "brief's conventions read: GBP %.2f per MWh, 11 kV nominal, resistances at 20 C" % VALUE)

    iis = re.sub(r"\s+", " ", pdf_text("iis_reporting_definitions_2025-26.pdf"))
    ev = iis.split("Exceptional event accepted for 2025/26", 1)[1]
    storm = [f for f in feeders if f in ev]
    ok(storm == ["Oxenford Lane", "Kestle", "Stowmere"] and "3 minutes or longer" in iis and "half weight" in iis,
       "reporting definitions read: storm circuits %s, 3-minute threshold, half weight for pre-arranged" % storm)
    S0, S1 = dt.datetime(2026, 1, 6, 15, 0), dt.datetime(2026, 1, 7, 23, 59)

    # ---------------------------------------------------------------- the network and its metering
    wb = openpyxl.load_workbook(os.path.join(T, "hv_network_model.xlsx"), read_only=True)
    secs = [dict(zip(["id", "circ", "a", "b", "m", "cond", "cons", "yr", "state"], r))
            for r in list(wb["Sections"].iter_rows(values_only=True))[1:]]
    nodes = {r[0]: r for r in list(wb["Nodes"].iter_rows(values_only=True))[1:]}
    ok(sum(s["state"] == "Open" for s in secs) == 4 and all(s["cond"] in R20 for s in secs),
       "network model read: %d sections, 4 normally open, every conductor in the data sheet" % len(secs))
    down = {}
    for s in secs:
        if s["state"] == "Closed":
            down.setdefault(s["a"], []).append(s["b"])

    def beyond(node):
        out, stack = [], [node]
        while stack:
            n = stack.pop()
            out.append(n)
            stack.extend(down.get(n, []))
        return out

    hvreg = {r["agreement_ref"]: r for r in csv.DictReader(open(os.path.join(T, "hv_connections.csv")))}
    start = dt.datetime(2025, 3, 31, 23, 0)
    N = 17520
    tb = pq.read_table(os.path.join(T, "lv_monitoring_2025-26.parquet")).to_pydict()
    mon = {}
    for site, t, kw, kvar, kva, pf in zip(tb["site_no"], tb["period_start_utc"], tb["kw"], tb["kvar"], tb["kva"], tb["pf"]):
        P, Q = mon.setdefault(site, (np.zeros(N), np.zeros(N)))
        i = utc_index(t.replace(tzinfo=None), start)
        if kw is None or (isinstance(kw, float) and np.isnan(kw)):
            P[i], Q[i] = kva * pf, kva * (1 - pf * pf) ** 0.5
        else:
            P[i], Q[i] = kw, kvar
    nreg = sum(1 for _ in csv.DictReader(open(os.path.join(T, "secondary_substations.csv"))))
    ok(len(mon) == nreg and len(tb["site_no"]) > 1_000_000,
       "monitoring read: %d monitors (every registered substation), %d rows" % (len(mon), len(tb["site_no"])))
    st = {}
    per_day = {}
    for r in csv.DictReader(open(os.path.join(T, "hv_settlement_2025-26.csv"))):
        d = dt.datetime.strptime(r["settlement_date"], "%Y%m%d")
        per_day[d] = max(per_day.get(d, 0), int(r["settlement_period"]))
        mid = d - dt.timedelta(hours=1 if bst(d - dt.timedelta(hours=1)) else 0)
        i = utc_index(mid + dt.timedelta(minutes=30 * (int(r["settlement_period"]) - 1)), start)
        A, Rq = st.setdefault(r["mpan"], (np.zeros(N), np.zeros(N)))
        A[i] = 2 * float(r["active_kwh"])
        Rq[i] = 2 * float(r["reactive_kvarh"])
    ok(per_day[dt.datetime(2025, 10, 26)] == 50 and per_day[dt.datetime(2026, 3, 29)] == 46 and len(per_day) == 365,
       "settlement periods: 50 on 26 October 2025, 46 on 29 March 2026, 365 settlement days")

    def flow(node):
        ref = nodes[node][3]
        if ref is None:
            return None
        ref = str(ref)
        if ref.startswith("CA-"):
            h = hvreg[ref]
            pi, qi = st[h["import_mpan"]]
            if h["export_mpan"]:
                pe, qe = st[h["export_mpan"]]
                return pi - pe, qi - qe
            return pi, qi
        return mon.get(int(ref.replace("SS", "").strip()))

    save = {}
    for f, (ids, new) in PROJ.items():
        tot = 0.0
        for sid in ids:
            s = [x for x in secs if x["id"] == sid][0]
            P, Q = np.zeros(N), np.zeros(N)
            for n in beyond(s["b"]):
                fl = flow(n)
                if fl is not None:
                    P, Q = P + fl[0], Q + fl[1]
            dR = (R20[s["cond"]] - R20[new]) * s["m"] / 1000
            tot += float(((P ** 2 + Q ** 2) * dR / 121.0 / 1000 * 0.5).sum())
        save[f] = tot / 1000
    order = sorted(save, key=lambda k: -save[k])

    # ---------------------------------------------------------------- the interruption measures
    rows = list(csv.DictReader(open(os.path.join(T, "interruptions_2023-2026.csv"))))
    cust = list(csv.DictReader(open(os.path.join(T, "customers_by_substation_2023-2026.csv"))))
    circ = {int(r["site_no"]): r["circuit"].replace(" (HV)", "") for r in cust}
    den = {}
    for r in cust:
        if r["month_end"] == "2026-03-31":
            den[circ[int(r["site_no"])]] = den.get(circ[int(r["site_no"])], 0) + int(r["customers"])
    hw = openpyxl.load_workbook(os.path.join(T, "generator_hire_2025-26.xlsx"), read_only=True)
    hires = [(int(r[4]), dt.datetime.strptime(r[8], "%d/%m/%Y %H:%M"), dt.datetime.strptime(r[9], "%d/%m/%Y %H:%M"))
             for r in list(hw["Hire log"].iter_rows(values_only=True))[1:]]
    inc = {}
    for r in rows:
        t = dt.datetime.fromisoformat(r["step_time"])
        inc.setdefault(r["reference"], dict(meta=r, ev={}))["ev"].setdefault(int(r["site_no"]), []).append(
            (t, int(r["customers_off"] or 0) - int(r["customers_restored"] or 0)))
    spans = {}
    for ref, x in inc.items():
        if x["meta"]["voltage"] != "LV":
            for s, ev in x["ev"].items():
                ev.sort()
                spans.setdefault(s, []).append((ev[0][0], min(t for t, n in ev if n < 0)))
    ci, cml = {f: 0.0 for f in feeders}, {f: 0.0 for f in feeders}
    lo, hi = dt.datetime(2025, 4, 1), dt.datetime(2026, 4, 1)
    dropped_dup = held = 0
    for ref, x in inc.items():
        m = x["meta"]
        t0 = dt.datetime.fromisoformat(m["incident_start"])
        if not (lo <= t0 < hi) or m["voltage"] == "EHV" or (S0 <= t0 <= S1 and m["circuit"] in storm):
            continue
        if m["voltage"] == "LV":
            s0 = next(iter(x["ev"]))
            a = min(t for t, n in x["ev"][s0])
            if any(p <= a < q for p, q in spans.get(s0, [])):
                dropped_dup += 1
                continue
        w = 0.5 if m["category"] == "Pre-arranged" else 1.0
        for s, ev in x["ev"].items():
            ev.sort()
            a, n = ev[0]
            b = min(t for t, k in ev if k < 0)
            g = [h for h in hires if h[0] == s and (h[1] <= a and h[2] >= b or a <= h[1] < b)]
            if g and g[0][1] <= a:
                held += 1
                continue
            if g:
                b = g[0][1]
            mins = (b - a).total_seconds() / 60
            if mins < 3:
                continue
            ci[circ[s]] += w * n
            cml[circ[s]] += w * n * mins
    IIS = {f: (round(100 * ci[f] / den[f], 1), round(cml[f] / den[f], 1)) for f in feeders}
    ok(dropped_dup > 0 and held > 0, "devices present: %d work-management jobs inside HV outages, %d sites held on generator"
       % (dropped_dup, held))

    # ---------------------------------------------------------------- against figures.json, the goldens and the write-up
    F = json.load(open(os.path.join(BASE, "generator", "figures.json")))
    ok(order[0] == F["call"] and order[1] == F["runner_up"], "call %s, runner-up %s" % (order[0], order[1]))
    for f in feeders:
        ok(round(save[f], 1) == F["projects"][f]["year_mwh"], "%s saving %.1f MWh" % (f, save[f]))
        ok(int(round(save[f] * VALUE, -2)) == F["projects"][f]["value_gbp"], "%s value GBP %d" % (f, round(save[f] * VALUE, -2)))
        ok(IIS[f] == (F["iis"][f]["ci"], F["iis"][f]["cml"]), "%s CI %.1f CML %.1f" % (f, IIS[f][0], IIS[f][1]))
    ok(round(save[order[0]], 1) - round(save[order[1]], 1) == F["gap_mwh"] or abs(round(save[order[0]] - save[order[1]], 1) - F["gap_mwh"]) < 1e-9,
       "gap %.1f MWh" % F["gap_mwh"])
    gw = openpyxl.load_workbook(os.path.join(G, "losses_programme_2027.xlsx"), read_only=True)
    grows = [r for r in gw["Programme"].iter_rows(min_row=5, values_only=True) if r[1] in feeders]
    ok(len(grows) == 8 and all(abs(r[7] - round(save[r[1]], 1)) < 1e-9 and r[8] == int(round(save[r[1]] * VALUE, -2)) for r in grows)
       and [r[0] for r in grows if r[0]] == ["Fund", "Runner-up"] and grows[0][1] == order[0],
       "golden workbook: eight rows agree, Fund and Runner-up marked")
    gc = list(csv.DictReader(open(os.path.join(G, "iis_feeders_2025-26.csv"))))
    ok(len(gc) == 8 and all((float(r["ci_per_100_customers"]), float(r["cml_per_customer"])) == IIS[r["feeder"]] for r in gc),
       "golden CSV: eight feeders' CI and CML agree")
    sub = open(os.path.join(BASE, "submission.md")).read()
    ok(all("%s: %.1f MWh" % (f, round(save[f], 1)) in sub and "%s: CI %.1f, CML %.1f" % (f, *IIS[f]) in sub for f in feeders),
       "submission.md block 4 carries every saving, CI and CML")
    print("%s  verify: %d failures" % ("PASS" if not FAIL else "FAIL", len(FAIL)))
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
