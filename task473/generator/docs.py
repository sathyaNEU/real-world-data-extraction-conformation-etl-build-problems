"""The pack's documents and workbooks: the GHSO's FY2028 task force guidelines, the census's reference manual, the
crash records system's extract notes, the laboratory's extract notes and its 2025 annual report tables, the agency
directory, the DOT's county miles and the motor vehicle division's licensed drivers."""
import datetime as dt

from reportlab import rl_config

rl_config.invariant = 1
from reportlab.lib import colors                                                      # noqa: E402
from reportlab.lib.pagesizes import letter                                            # noqa: E402
from reportlab.lib.styles import ParagraphStyle                                       # noqa: E402
from reportlab.lib.units import inch                                                  # noqa: E402
from reportlab.pdfbase import pdfmetrics                                              # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont                                          # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

import agencies as AG   # noqa: E402
import geo as G         # noqa: E402
import lab as L         # noqa: E402
import truth as T       # noqa: E402
import vmt as V         # noqa: E402
import world as W       # noqa: E402

FONT_DIR = "/usr/share/fonts/truetype/liberation/"
for nm, f in (("LiberationSerif", "Regular"), ("LiberationSerif-Bold", "Bold"), ("LiberationSerif-Italic", "Italic"),
              ("LiberationSans", "Regular"), ("LiberationSans-Bold", "Bold"), ("LiberationSans-Italic", "Italic")):
    try:
        pdfmetrics.getFont(nm)
    except KeyError:
        pdfmetrics.registerFont(TTFont(nm, FONT_DIR + nm.split("-")[0] + "-" + f + ".ttf"))
pdfmetrics.registerFontFamily("LiberationSerif", normal="LiberationSerif", bold="LiberationSerif-Bold",
                              italic="LiberationSerif-Italic", boldItalic="LiberationSerif-Bold")
pdfmetrics.registerFontFamily("LiberationSans", normal="LiberationSans", bold="LiberationSans-Bold",
                              italic="LiberationSans-Italic", boldItalic="LiberationSans-Bold")
INK = colors.HexColor("#1d1f22")
GREY = colors.HexColor("#5b6067")
NAVY = colors.HexColor("#1f3a5f")
MAROON = colors.HexColor("#6b1d2a")


def styles(font="LiberationSans", accent=NAVY, size=9.6):
    bold = font + "-Bold"
    return dict(
        BODY=ParagraphStyle("body", fontName=font, fontSize=size, leading=size * 1.36, textColor=INK, spaceAfter=5),
        CELL=ParagraphStyle("cell", fontName=font, fontSize=7.9, leading=9.8, textColor=INK),
        CELLB=ParagraphStyle("cellb", fontName=bold, fontSize=7.9, leading=9.8, textColor=INK),
        H1=ParagraphStyle("h1", fontName=bold, fontSize=13.5, leading=17, textColor=accent, spaceAfter=2),
        H1S=ParagraphStyle("h1s", fontName=font, fontSize=8.8, leading=11.5, textColor=GREY, spaceAfter=9),
        H2=ParagraphStyle("h2", fontName=bold, fontSize=10.2, leading=13.5, textColor=INK, spaceBefore=7,
                          spaceAfter=3))


S = styles()
SERIF = styles("LiberationSerif", MAROON, 10.2)


def _doc(path, title, footer, author):
    def frame(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(GREY)
        canvas.setFont("LiberationSans", 7.3)
        canvas.drawString(0.85 * inch, 0.58 * inch, footer)
        canvas.drawRightString(letter[0] - 0.85 * inch, 0.58 * inch, "Page %d" % doc.page)
        canvas.restoreState()
    d = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.85 * inch, rightMargin=0.85 * inch,
                          topMargin=0.78 * inch, bottomMargin=0.88 * inch, title=title, author=author, subject="",
                          creator="", producer="")
    return d, frame


def _table(rows, widths, st=S, head="#dfe6ef"):
    data = [[Paragraph(str(v), st["CELLB"] if i == 0 else st["CELL"]) for v in r] for i, r in enumerate(rows)]
    t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa1a9")),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(head)),
                           ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 3.5),
                           ("RIGHTPADDING", (0, 0), (-1, -1), 3.5)]))
    return t


def _story(title, sub, sections, st=S):
    story = [Paragraph(title, st["H1"]), Paragraph(sub, st["H1S"])]
    for head, paras in sections:
        if head:
            story.append(Paragraph(head, st["H2"]))
        for p in paras:
            story.append(p if not isinstance(p, str) else Paragraph(p, st["BODY"]))
    return story


# ------------------------------------------------------------------------------------------------ the guidelines
GUIDELINE_DATE = "9 February 2027"
SECTIONS = {"estimand": "2.2", "drivers": "2.3", "rate": "2.4", "baselines": "3.1", "testing": "4.2"}
WINDOW_DAYS = 30


def write_guidelines(path):
    d, frame = _doc(path, "FY2028 County Impaired Driving Task Force Grants: Program Guidelines",
                    "%s  |  FY2028 County Impaired Driving Task Force Grants, Program Guidelines" % W.GHSO, W.GHSO)
    secs = [
        ("1. Purpose", [
            "1.1 The FY2028 County Impaired Driving Task Force grants fund a task force in each county where impaired "
            "driving kills people at a high rate: sustained enforcement by the county's agencies together, a "
            "prosecutor and a probation officer for impaired-driving cases, a treatment court referral line and a "
            "public information campaign. Task forces run from 1 October 2027 to 30 September 2028 and are funded "
            "from the GHSO's federal impaired-driving countermeasures funds.",
            "1.2 Every high-range county (section 2) may apply. A county that is not high-range may apply for "
            "impaired-driving funds only under the GHSO's general program."]),
        ("2. High-range counties", [
            "2.1 The %s, which keeps the state's fatal crash census under its agreement with the GHSO, certifies to "
            "the GHSO by 1 May 2027 the list of high-range counties for FY2028 and each county's rate under 2.4, in a "
            "letter with a chart of the 14 counties' rates." % W.CENTER,
            "%s A county's alcohol-impaired-driving fatalities are the people who died within %d days of a crash on "
            "a public road in the county that occurred in 2023, 2024 or 2025 in which at least one driver had a blood "
            "alcohol concentration (BAC) of .08 g/dL or more, whether or not that driver was tested. Where a "
            "driver's BAC is not known, the count is the number expected from the records: each crash counts its "
            "deaths times the chance that at least one of its drivers was at .08 g/dL or more." % (
                SECTIONS["estimand"], WINDOW_DAYS),
            "%s A driver is the operator of a motor vehicle in transport, a motorcycle rider included. Passengers, "
            "pedestrians and pedalcyclists are not drivers, and a driver who left the scene and was never identified "
            "is not counted as a driver." % SECTIONS["drivers"],
            "%s A county's rate is its 2023 to 2025 alcohol-impaired-driving fatalities divided by its 2023 to 2025 "
            "vehicle miles traveled, per 100 million vehicle miles, using the %s's annual county vehicle miles. A "
            "county whose rate is %.2f or more is high-range." % (SECTIONS["rate"], W.DOT, W.THRESHOLD)]),
        ("3. Applications", [
            "%s Each application states its county's 2025 baselines, which the Center certifies with the list:" %
            SECTIONS["baselines"],
            "(a) the impaired-driving arrests made in the county in 2025 for which the %s received a blood specimen "
            "(an arrest's date is the date in Central time its first specimen was collected), and the per cent of "
            "all those arrests, a specimen tested only for drugs included, in which the laboratory's whole-blood "
            "ethanol result for the first specimen collected was .15 g/dL or more; and" % W.LAB,
            "(b) the injury crashes on public roads in the county in 2025 in which no one died within 30 days and the "
            "investigating officer reported that a driver had been drinking; an injury crash is one in which someone "
            "was injured, a possible injury included.",
            "3.2 An application names the county's lead agency, the participating agencies and the prosecutor's "
            "office, and states the task force's FY2028 targets against the 2025 baselines. Applications are due "
            "15 June 2027."]),
        ("4. Timeline", [
            _table([["Date", "Step"],
                    ["1 May 2027", "The Center certifies the FY2028 high-range counties, their rates and their 2025 "
                                   "baselines (2.1, 3.1)"],
                    ["15 June 2027", "Task force applications due"],
                    ["1 July 2027", "FY2028 Highway Safety Plan submitted"],
                    ["15 September 2027", "Awards"],
                    ["1 October 2027", "Task forces begin"]],
                   [1.35 * inch, 5.3 * inch])]),
    ]
    story = _story("FY2028 County Impaired Driving Task Force Grants: Program Guidelines",
                   "%s. Issued %s. Contact: %s, %s." % (W.GHSO, GUIDELINE_DATE, W.PEOPLE["ghso"][0],
                                                       W.PEOPLE["ghso"][1]), secs)
    d.build(story, onFirstPage=frame, onLaterPages=frame)


# ------------------------------------------------------------------------------------------------ the census manual
def write_manual(path):
    d, frame = _doc(path, "Fatal Crash Census Reference Manual, 2024 edition",
                    "%s  |  Fatal Crash Census Reference Manual, 2024 edition" % W.CENTER, W.CENTER)
    old_fields = [["Field", "Contents and codes (2011 report form, crash years through 2023)"],
                  ["CASE_NUM", "Census case number"],
                  ["ACC_DATE, ACC_TIME", "Crash date (MM/DD/YYYY) and time (HHMM, 24-hour clock, local time; 9999 "
                                         "unknown)"],
                  ["CNTY, CITY_NAME", "County FIPS code where the crash occurred; city if inside one"],
                  ["AGCY, RPT_NUM", "Investigating agency (four-digit code) and its crash report number"],
                  ["RTE_SYS, RTE_ID", "1 Interstate, 2 US route, 3 state route, 4 county road, 5 city street; "
                                      "route or street"],
                  ["LGT, WTHR, COLL", "Light (1 daylight, 2 dark not lighted, 3 dark lighted, 4 dawn, 5 dusk); "
                                      "weather (1 clear, 2 rain, 3 snow, 4 fog, 5 cloudy); manner of collision "
                                      "(0 not a collision with a motor vehicle, 1 front to rear, 2 head on, 3 angle, "
                                      "4 sideswipe)"],
                  ["VEH_CT, FATAL_CT, LAT, LON", "Motor vehicles; people who died within 30 days; location"],
                  ["VEH_NUM, PER_NUM", "Vehicle number (0 for a pedestrian or pedalcyclist) and person number"],
                  ["PER_TYP", "1 driver, 2 passenger, 5 pedestrian, 6 pedalcyclist"],
                  ["BODY", "01 passenger car, 02 sport utility vehicle, 03 pickup, 04 van, 05 motorcycle, 06 truck, "
                           "00 not a motor vehicle"],
                  ["AGE, SEX", "Age in years (999 unknown); 1 male, 2 female, 9 unknown"],
                  ["INJ", "0 no injury, 1 possible, 2 suspected minor, 3 suspected serious, 4 fatal (died within 30 "
                          "days), 9 unknown"],
                  ["DTH_DATE", "Date of death"],
                  ["HIT_RUN", "1 the driver left the scene and was not identified"],
                  ["REST", "Restraint or helmet: 03 lap and shoulder belt, 05 helmet, 20 none used, 99 unknown, 00 "
                           "not applicable"],
                  ["ALC_TST", "Alcohol test: 0 not tested, 1 blood, 3 vitreous humor"],
                  ["ALC_RES", "Alcohol test result in thousandths of a gram per deciliter (085 is .085 g/dL); 996 "
                              "not tested"],
                  ["DWI_ARR", "Driver arrested at the scene for driving while impaired: 1 yes, 0 no"]]
    new_fields = [["Field (2024 report form, crash years from 2024)", "Contents"],
                  ["crash_id", "Census case number"],
                  ["crash_date, crash_time", "Crash date and time (HH:MM, local time; blank unknown)"],
                  ["county_fips, municipality", "County where the crash occurred; city if inside one"],
                  ["agency_id, report_number", "Investigating agency (agency number) and its crash report number"],
                  ["road_class, road_name, light_condition, weather, collision_type", "As the 2011 form's codes, "
                                                                                       "in words"],
                  ["motor_vehicles, fatalities, latitude, longitude", "As VEH_CT, FATAL_CT, LAT, LON"],
                  ["unit, person", "Unit and person number; a pedestrian or pedalcyclist is a unit of its own"],
                  ["person_type, unit_type", "DRIVER, PASSENGER, PEDESTRIAN, PEDALCYCLIST; the unit's body type "
                                             "(NON-MOTORIST for a pedestrian or pedalcyclist)"],
                  ["age, sex, injury", "Age; M, F or U; K (died within 30 days), A, B, C, O or U"],
                  ["date_of_death, hit_and_run, restraint", "As DTH_DATE, HIT_RUN (Y or N) and REST"],
                  ["alcohol_test, alcohol_result", "BLOOD, VITREOUS or NOT TESTED; the result in g/dL"],
                  ["arrested_at_scene", "Y or N, drivers only"]]
    secs = [
        ("1. The census", [
            "1.1 The Fatal Crash Census holds every crash on a public road in the state in which a person died within "
            "30 days of the crash. The Center compiles it from the investigating agency's crash report, the coroners' "
            "and the State Medical Examiner's reports, the %s's results and death certificates." % W.LAB,
            "1.2 A year's file is closed when the 30-day window of its last crash has passed and the laboratory has "
            "reported on every specimen, and is not changed after it closes."]),
        ("2. Files and layouts", [
            "2.1 Each year has a crash file (one row per crash) and a person file (one row per person in the crash, "
            "every vehicle occupant and every pedestrian or pedalcyclist involved). Crash years through 2023 are in "
            "the layout of the 2011 crash report form; from 2024 the census follows the 2024 report form, which "
            "replaced codes with words, gave each agency a new agency number and records alcohol results in g/dL. "
            "The agency directory carries each agency's former code beside its agency number.",
            _table(old_fields, [1.75 * inch, 4.9 * inch]), Spacer(1, 6),
            _table(new_fields, [2.6 * inch, 4.05 * inch])]),
        ("3. Injuries and deaths", [
            "3.1 A person coded fatal (INJ 4, injury K) died within 30 days of the crash, and the date of death is "
            "recorded. A person who died more than 30 days after the crash keeps the injury the officer reported, "
            "with the date of death recorded; that death is not a traffic fatality."]),
        ("4. Alcohol tests", [
            "4.1 A test result is the laboratory's ethanol result in g/dL to three decimals as the laboratory reports "
            "it (truncated): whole blood, or vitreous humor for some killed persons.",
            "%s Coroners and the State Medical Examiner submit blood or vitreous humor from every driver killed and "
            "from most other persons killed. The State Patrol's fatal crash procedure draws blood from every "
            "surviving driver, and the Corwin County and Ossian County sheriffs follow the same procedure. Other "
            "sheriffs' offices and police departments draw blood from a surviving driver arrested at the scene for "
            "impaired driving, with a warrant when the driver refuses. A surviving driver from whom no blood was "
            "drawn is recorded as not tested." % SECTIONS["testing"],
            "4.3 A driver who left the scene and was not identified is recorded with the vehicle, with no age, sex, "
            "injury or test."]),
        ("Edition", [
            "2024 edition, issued 15 March 2024 with the 2024 report form, replacing the 2019 edition. Questions to "
            "%s, %s." % (W.PEOPLE["records"][0], W.PEOPLE["records"][1])]),
    ]
    story = _story("Fatal Crash Census Reference Manual", "%s. 2024 edition." % W.CENTER, secs)
    d.build(story, onFirstPage=frame, onLaterPages=frame)


# ------------------------------------------------------------------------------------------------ crash records notes
def write_crs_notes(path, n_reports, n_units):
    d, frame = _doc(path, "Crash Records System: 2025 injury and fatal crash report extract",
                    "%s  |  Crash Records System extract notes, 2025 reports" % W.CENTER, W.CENTER)
    secs = [
        ("Coverage", [
            "The extract holds every version of every crash report an investigating agency filed with the Crash "
            "Records System for a crash in 2025 in which the officer reported a person killed or injured, as the "
            "system stood on 8 March 2027: %s report versions and %s units." % (format(n_reports, ","),
                                                                                format(n_units, ","))]),
        ("Files", [
            _table([["File", "One row per", "Fields"],
                    [W.FILES["reports"], "report version", "report_id, version, version_status, submitted, agency_id, "
                     "report_number, crash_date, crash_time (local), county_fips, municipality, road_class, "
                     "road_name, latitude, longitude, public_road, severity, motor_vehicles, non_motorists"],
                    [W.FILES["units"], "unit in a report version", "report_id, version, unit, unit_type, driver_age, "
                     "driver_sex, driver_condition, unit_severity"]],
                   [1.6 * inch, 1.15 * inch, 3.9 * inch])]),
        ("Reports and versions", [
            "report_id is the system's number for a report; report_number is the investigating agency's. Version 0 is "
            "the report as first filed. A supplement (version 1 or later) replaces the report's fields and its units. "
            "The records unit rejects a supplement it cannot accept, and a rejected supplement is kept with "
            "version_status REJECTED and does not change the report. A report as it stands is its latest accepted "
            "version.",
            "severity is the most severe injury the officer reported in the crash at the time of the report, on the "
            "KABCO scale (K killed, A suspected serious injury, B suspected minor injury, C possible injury, O no "
            "injury). Deaths after a report are recorded in the Fatal Crash Census, not in the report.",
            "public_road is N for a crash on private property (a parking lot, a driveway) that an agency reported."]),
        ("Units", [
            "unit_type is MOTOR VEHICLE (a vehicle in transport), PEDESTRIAN or PEDALCYCLIST. driver_condition is "
            "the officer's observation of the unit's driver, or of the pedestrian or pedalcyclist: APPARENTLY NORMAL, "
            "HAD BEEN DRINKING, UNDER THE INFLUENCE OF DRUGS, ILLNESS, FATIGUED/ASLEEP or UNKNOWN. unit_severity "
            "is the most severe injury in the unit. driver_age and driver_sex are blank where the officer did not "
            "record them."]),
        ("Contact", ["%s, %s." % (W.PEOPLE["records"][0], W.PEOPLE["records"][1])]),
    ]
    story = _story("Crash Records System: 2025 injury and fatal crash report extract",
                   "%s, crash records unit. Extract of 8 March 2027." % W.CENTER, secs)
    d.build(story, onFirstPage=frame, onLaterPages=frame)


# ------------------------------------------------------------------------------------------------ the laboratory
def lab_notes_text(n_rows):
    lines = [
        "STATE TOXICOLOGY LABORATORY",
        "Specimen extract prepared for the %s" % W.CENTER,
        "Extract date: 12 March 2027",
        "",
        "Coverage: every specimen kit received from 1 January 2023 through 15 January 2026 (received date), %s rows."
        % format(n_rows, ","),
        "",
        "One row per analysis per tube.",
        "",
        "lab_number          Laboratory number of the specimen kit (T + year received + sequence). A kit holds",
        "                    two tubes; tube 2 is analyzed only on request or for confirmation.",
        "case_type           DUI: impaired-driving arrest. FATAL CRASH: blood drawn under the State Patrol's",
        "                    fatal crash procedure from a driver not arrested. DEATH INVESTIGATION: coroner or",
        "                    medical examiner. OTHER CRIMINAL: any other criminal case.",
        "submitting_agency   Agency number on the request form.",
        "agency_case_number  The submitting agency's case number.",
        "collection_site     Where the specimen was collected (site codes in %s)." % W.FILES["sites"],
        "collected_utc       Collection date and time from the chain-of-custody form, in UTC.",
        "received_date       Date the laboratory received the kit.",
        "specimen            WB whole blood, SER serum, VH vitreous humor, UR urine.",
        "analyte             ETHANOL or DRUG SCREEN (immunoassay screen; confirmations are reported separately",
        "                    and are not in this extract).",
        "result              ETHANOL: g/100 mL (g/dL), three decimals, truncated. Serum results are reported as",
        "                    measured; the whole-blood equivalent is the serum result divided by %.2f." % L.SERUM,
        "                    DRUG SCREEN: POSITIVE or NEGATIVE.",
        "report_status       FINAL. AMENDED: a result reissued after review, which replaces the earlier report;",
        "                    the earlier report is then marked SUPERSEDED. CANCELLED: request cancelled by the",
        "                    agency before analysis.",
        "report_date         Date of the report.",
        "",
        "An agency may submit a second specimen drawn later from the same person; it has its own lab number and",
        "the same agency case number.",
        "",
        "Questions: Records Section, State Toxicology Laboratory.",
    ]
    return "\r\n".join(lines) + "\r\n"


def write_lab_notes(path, n_rows):
    with open(path, "w", newline="") as fh:
        fh.write(lab_notes_text(n_rows))


def lab_annual(kits):
    """The laboratory's 2025 statistical tables: kits received in 2025 by case type, and the ethanol results of the
    DUI kits received in 2025 by range (first reports as issued)."""
    rec = [k for k in kits if k["received"].year == 2025]
    by_type = {}
    for k in rec:
        by_type[k["case_type"]] = by_type.get(k["case_type"], 0) + 1
    ranges = [("Below .01", 0.0, 0.010), (".01 to .079", 0.010, 0.080), (".08 to .149", 0.080, 0.150),
              (".15 to .199", 0.150, 0.200), (".20 or more", 0.200, 9.0)]
    dist = {r[0]: 0 for r in ranges}
    nd = 0
    for k in rec:
        if k["case_type"] != "DUI" or k["drug_only"] or k["cancel"]:
            continue
        v = k["bac"] + (k["re"] or 0.0)
        v = max(0.0, v)
        if k["specimen"] == "SER":
            v = L._trunc3(v * L.SERUM)
        nd += 1
        for name, lo, hi in ranges:
            if lo <= v < hi:
                dist[name] += 1
    return by_type, dist, nd, ranges


def write_lab_annual(path, kits):
    by_type, dist, nd, ranges = lab_annual(kits)
    d, frame = _doc(path, "State Toxicology Laboratory Annual Report 2025: Statistical Tables",
                    "%s  |  Annual Report 2025, statistical tables" % W.LAB, W.LAB)
    t1 = [["Case type", "Kits received in 2025"]] + [[k, format(by_type.get(k, 0), ",")] for k in
                                                       ("DUI", "FATAL CRASH", "DEATH INVESTIGATION", "OTHER CRIMINAL")]
    t1.append(["All case types", format(sum(by_type.values()), ",")])
    t2 = [["Ethanol, g/100 mL", "DUI kits", "Per cent"]] + [
        [name, format(dist[name], ","), "%.1f" % (100.0 * dist[name] / nd)] for name, lo, hi in ranges]
    t2.append(["All DUI kits with an ethanol result", format(nd, ","), "100.0"])
    secs = [
        ("About these tables", [
            "The tables count specimen kits by the date the laboratory received them, so a person from whom two "
            "specimens were drawn is counted twice, and a kit drawn late in December may be counted in the next "
            "year. Ethanol results are the first report issued on each kit, as measured (serum not converted)."]),
        ("Table 1. Kits received, 2025", [_table(t1, [3.2 * inch, 1.8 * inch])]),
        ("Table 2. Ethanol results of DUI kits received in 2025", [_table(t2, [3.2 * inch, 1.3 * inch, 1.0 * inch])]),
        ("Table 3. Turnaround", [
            "Median days from receipt to report in 2025: ethanol 9, drug screen 16. The laboratory's goal is 15 days "
            "for ethanol and 30 days for drug screens."]),
    ]
    story = _story("Annual Report 2025: Statistical Tables", "%s. Published 6 March 2027." % W.LAB, secs)
    d.build(story, onFirstPage=frame, onLaterPages=frame)


# ------------------------------------------------------------------------------------------------ workbooks
def _head(ws, row, fill="DCE6F1"):
    from openpyxl.styles import Font, PatternFill
    for c in ws[row]:
        c.font = Font(name="Calibri", bold=True)
        c.fill = PatternFill("solid", fgColor=fill)


def directory_rows():
    out = []
    for a in AG.AGENCIES:
        out.append([a.aid, a.old, a.name, a.kind, a.fips, W.NAMES[a.fips], a.city or ""])
    hq = W.PATROL_HQ
    out.append([AG.PATROL_ID, "9000", "State Patrol Headquarters", "State Patrol", hq, W.NAMES[hq], "Kessington"])
    for c in W.COUNTIES:
        f = c["fips"]
        if f not in L.ME_COUNTIES:
            out.append([L.CORONER_ID[f], "", "%s County Coroner" % c["name"], "Coroner", f, c["name"],
                        c["cities"][0]])
    out.append([L.ME_ID, "", "Office of the State Medical Examiner", "Medical Examiner", "013", W.NAMES["013"],
                "Kessington"])
    out.sort(key=lambda r: (r[3] != "State Patrol", r[4], r[0]))
    return out


def write_directory(path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Agencies"
    ws.append(["agency_id", "former_code", "agency_name", "agency_type", "county_fips", "county", "city"])
    _head(ws, 1)
    for r in directory_rows():
        ws.append([None if v == "" else v for v in r])
    for col, w in zip("ABCDEFG", (11, 12, 50, 18, 11, 12, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("State Patrol posts")
    ws2.append(["post", "agency_id", "post_name", "post_city", "county_fips", "county"])
    _head(ws2, 1)
    for c in W.COUNTIES:
        a = AG.patrol_post(c["fips"])
        ws2.append([a.post, a.aid, a.name, c["cities"][0], c["fips"], c["name"]])
    for col, w in zip("ABCDEF", (6, 11, 22, 16, 11, 12)):
        ws2.column_dimensions[col].width = w
    ws3 = wb.create_sheet("Notes")
    for line in ["Agency directory kept by the %s for the Fatal Crash Census and the Crash Records System." % W.CENTER,
                 "agency_id: the agency number adopted with the 2024 crash report form (1 January 2024). former_code: "
                 "the four-digit agency code on the 2011 report form.",
                 "Each State Patrol post patrols the county shown on the State Patrol posts sheet.",
                 "Updated 4 January 2027."]:
        ws3.append([line])
    ws3.column_dimensions["A"].width = 112
    wb.save(path)


def write_dvmt(path, tab):
    from openpyxl import Workbook
    wb = Workbook()
    first = True
    for y in W.YEARS:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = "DVMT %d" % y
        ws.append(["County FIPS", "County"] + list(V.CLASSES) + ["All classes"])
        _head(ws, 1, fill="E2EFDA")
        for c in W.COUNTIES:
            f = c["fips"]
            vals = [tab[f][(y, cl)] for cl in V.CLASSES]
            ws.append([f, c["name"]] + vals + [round(sum(vals), 1)])
        tot = [round(sum(tab[c["fips"]][(y, cl)] for c in W.COUNTIES), 1) for cl in V.CLASSES]
        ws.append([None, "State"] + tot + [round(sum(tot), 1)])
        for row in ws.iter_rows(min_row=2, min_col=3, max_col=3 + len(V.CLASSES)):
            for cell in row:
                cell.number_format = "#,##0.0"
        for col, w in zip("ABCDEFGHIJ", (11, 12, 11, 15, 15, 13, 13, 13, 11, 12)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "C2"
    ws = wb.create_sheet("Notes")
    for line in ["%s, Transportation Planning Division: county vehicle miles traveled, 2023 to 2025." % W.DOT,
                 "Values are annual average daily vehicle miles traveled (DVMT) on public roads, in thousands of "
                 "vehicle miles, by functional class, from the state's HPMS submissions.",
                 "Annual vehicle miles traveled are DVMT times the number of days in the year.",
                 "2023 and 2024 are final; 2025 is final as of 30 January 2027."]:
        ws.append([line])
    ws.column_dimensions["A"].width = 110
    wb.save(path)


def licensed_rows(vmt):
    out = []
    for c in W.COUNTIES:
        f = c["fips"]
        rng = T.rng_for("licensed", f)
        v25 = vmt[f] * V.YEAR_SPLIT[2025]
        total = int(round(v25 / float(rng.uniform(12400, 14600))))
        u21 = int(round(total * float(rng.uniform(0.045, 0.062) if f != "027" else 0.094)))
        o65 = int(round(total * float(rng.uniform(0.19, 0.26))))
        out.append([f, c["name"], total, u21, total - u21 - o65, o65])
    return out


def write_licensed(path, vmt):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Licensed drivers 2025"
    ws.append(["County FIPS", "County", "Licensed drivers", "Under 21", "21 to 64", "65 and over"])
    _head(ws, 1, fill="FCE4D6")
    for r in licensed_rows(vmt):
        ws.append(r)
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=6):
        for cell in row:
            cell.number_format = "#,##0"
    for col, w in zip("ABCDEF", (11, 12, 16, 11, 11, 12)):
        ws.column_dimensions[col].width = w
    ws2 = wb.create_sheet("Notes")
    for line in ["Department of Revenue, Motor Vehicle Division: licensed drivers by county of residence, 31 December "
                 "2025, from the driver license file.",
                 "Includes Class A, B, C and M licenses and instruction permits; excludes identification cards."]:
        ws2.append([line])
    ws2.column_dimensions["A"].width = 100
    wb.save(path)
