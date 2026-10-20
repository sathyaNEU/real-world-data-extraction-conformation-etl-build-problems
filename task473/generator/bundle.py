"""Reading the cut pack the way a solver must: every file through its own format, no generator state. The census in
both layouts, the agency directory, the DOT's miles, the laboratory's extract, the crash records extract, the
documents' text. verify.py and checks.py read the bundle only through here."""
import csv
import datetime as dt
import os
import subprocess
import zipfile
from zoneinfo import ZoneInfo

CENTRAL = ZoneInfo("America/Chicago")
NAMES = {}


def path(target, name):
    return os.path.join(target, name)


def find(target, prefix):
    hits = [f for f in os.listdir(target) if f.startswith(prefix)]
    assert len(hits) == 1, (prefix, hits)
    return os.path.join(target, hits[0])


def read_csv(p):
    with open(p, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def pdf_text(p):
    return subprocess.run(["pdftotext", "-layout", p, "-"], capture_output=True, text=True, check=True).stdout


def xlsx_sheets(p):
    from openpyxl import load_workbook
    wb = load_workbook(p, read_only=True, data_only=True)
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


# ------------------------------------------------------------------------------------------------- the directory
def directory(target):
    sh = xlsx_sheets(find(target, "fcc_agency_directory"))
    rows = sh["Agencies"]
    head = rows[0]
    ag = [dict(zip(head, r)) for r in rows[1:]]
    posts = [dict(zip(sh["State Patrol posts"][0], r)) for r in sh["State Patrol posts"][1:]]
    by_id = {a["agency_id"]: a for a in ag}
    by_old = {str(a["former_code"]): a for a in ag if a["former_code"] not in (None, "")}
    counties = {}
    for a in ag:
        counties[str(a["county_fips"])] = a["county"]
    return dict(agencies=by_id, by_old=by_old, posts={int(p["post"]): p for p in posts}, counties=counties)


def complete_agencies(dr):
    """Manual 4.2: the State Patrol (its posts) and the Corwin County and Ossian County sheriffs."""
    out = set()
    posts = {p["agency_id"] for p in dr["posts"].values()}
    for aid, a in dr["agencies"].items():
        if a["agency_type"] == "State Patrol" and aid in posts:
            out.add(aid)
        if a["agency_type"] == "Sheriff" and a["county"] in ("Corwin", "Ossian"):
            out.add(aid)
    return out


# ------------------------------------------------------------------------------------------------- the census
def census(target, dr=None):
    """Every 2023 to 2025 crash as {case, year, date, fips, agency, report, deaths, persons: [...]}, both layouts
    read into one structure; agencies on the 2024 numbers."""
    dr = dr or directory(target)
    out = []
    c23 = read_csv(find(target, "fcc_2023_crash"))
    p23 = read_csv(find(target, "fcc_2023_person"))
    per = {}
    for r in p23:
        per.setdefault(r["CASE_NUM"], []).append(r)
    for r in c23:
        m, d, y = (int(v) for v in r["ACC_DATE"].split("/"))
        date = dt.date(y, m, d)
        ag = dr["by_old"][r["AGCY"]]["agency_id"]
        ps = []
        for q in per.get(r["CASE_NUM"], []):
            role = {"1": "driver", "2": "passenger", "5": "pedestrian", "6": "pedalcyclist"}[q["PER_TYP"]]
            dd = None
            if q["DTH_DATE"]:
                mm, dd_, yy = (int(v) for v in q["DTH_DATE"].split("/"))
                dd = dt.date(yy, mm, dd_)
            tested = q["ALC_TST"] in ("1", "3")
            ps.append(dict(role=role, unit=int(q["VEH_NUM"]), killed=q["INJ"] == "4", death=dd,
                           identified=not (role == "driver" and q["HIT_RUN"] == "1"), tested=tested,
                           test=q["ALC_TST"], bac=int(q["ALC_RES"]) / 1000.0 if tested else None,
                           arrested=q["DWI_ARR"] == "1", body=q["BODY"]))
        out.append(dict(case=r["CASE_NUM"], year=y, date=date, fips=r["CNTY"], agency=ag, report=r["RPT_NUM"],
                        deaths=sum(1 for p in ps if p["killed"]), persons=ps, time=r["ACC_TIME"]))
    c25 = read_csv(find(target, "fcc_2024_2025_crash"))
    p25 = read_csv(find(target, "fcc_2024_2025_person"))
    per = {}
    for r in p25:
        per.setdefault(r["crash_id"], []).append(r)
    for r in c25:
        date = dt.date.fromisoformat(r["crash_date"])
        ps = []
        for q in per.get(r["crash_id"], []):
            role = q["person_type"].lower()
            tested = q["alcohol_test"] in ("BLOOD", "VITREOUS")
            ps.append(dict(role=role, unit=int(q["unit"]), killed=q["injury"] == "K",
                           death=dt.date.fromisoformat(q["date_of_death"]) if q["date_of_death"] else None,
                           identified=not (role == "driver" and q["hit_and_run"] == "Y"), tested=tested,
                           test=q["alcohol_test"], bac=float(q["alcohol_result"]) if tested else None,
                           arrested=q["arrested_at_scene"] == "Y", body=q["unit_type"]))
        out.append(dict(case=r["crash_id"], year=date.year, date=date, fips=r["county_fips"], agency=r["agency_id"],
                        report=r["report_number"], deaths=sum(1 for p in ps if p["killed"]), persons=ps,
                        time=r["crash_time"], road=r["road_name"]))
    return out


def observed(crashes, complete, roles=("driver",), unidentified=False, late_deaths=False):
    """The structure methods.py reads: each crash's drivers with killed, identified, tested, bac, arrested and
    complete. The keywords read the census another way: other person types as drivers, the unidentified drivers
    as untested ones, a death recorded after 30 days as a death."""
    obs = []
    for cr in crashes:
        ds = []
        for p in cr["persons"]:
            if p["role"] not in roles:
                continue
            ident = p["identified"] or unidentified
            ds.append(dict(killed=p["killed"], identified=ident, tested=p["tested"] and p["identified"],
                           bac=p["bac"], arrested=p["arrested"],
                           complete=(not p["killed"]) and cr["agency"] in complete,
                           agency="SP" if cr["agency"].startswith("SP") else cr["agency"]))
        deaths = cr["deaths"] + (sum(1 for p in cr["persons"] if p["death"] and not p["killed"]) if late_deaths else 0)
        obs.append(dict(fips=cr["fips"], year=cr["year"], deaths=deaths, drivers=ds, cid=cr["case"]))
    return obs


# ------------------------------------------------------------------------------------------------- the miles
def vmt(target, days=None):
    """Each county's 2023 to 2025 vehicle miles: each year's DVMT (thousands a day, all classes) times its days."""
    days = days or {2023: 365, 2024: 366, 2025: 365}
    sh = xlsx_sheets(find(target, "dot_county_dvmt"))
    out = {}
    for y in (2023, 2024, 2025):
        rows = sh["DVMT %d" % y]
        head = rows[0]
        for r in rows[1:]:
            rec = dict(zip(head, r))
            if not rec["County FIPS"]:
                continue
            tot = sum(float(rec[h]) for h in head[2:-1])
            out[str(rec["County FIPS"])] = out.get(str(rec["County FIPS"]), 0.0) + tot * 1000.0 * days[y]
    return out


def golden_g(obs, miles):
    """The golden completion on its own code path: the catch rate at each band where every surviving driver is
    tested, each arrest-only positive counted as one over it, the excess spread over the county's untested
    surviving drivers, each crash's deaths times the chance a driver was at .08 or more."""
    rows = [d for cr in obs for d in cr["drivers"] if d["complete"] and d["tested"] and not d["killed"]]
    catch = {}
    for b, lo, hi in ((0, 0.080, 0.150), (1, 0.150, 9.9)):
        r = [d for d in rows if lo <= d["bac"] < hi]
        catch[b] = sum(1 for d in r if d["arrested"]) / float(len(r))
    missed, untested = {}, {}
    for cr in obs:
        for d in cr["drivers"]:
            if d["killed"] or d["complete"] or not d["identified"]:
                continue
            if d["tested"] and d["bac"] >= 0.080:
                b = 0 if d["bac"] < 0.150 else 1
                missed[cr["fips"]] = missed.get(cr["fips"], 0.0) + 1.0 / catch[b] - 1.0
            elif not d["tested"]:
                untested[cr["fips"]] = untested.get(cr["fips"], 0) + 1
    q = {f: missed.get(f, 0.0) / n for f, n in untested.items()}
    deaths = {}
    for cr in obs:
        ds = [d for d in cr["drivers"] if d["identified"]]
        if any(d["tested"] and d["bac"] >= 0.080 for d in ds):
            p = 1.0
        else:
            keep = 1.0
            for d in ds:
                if not d["tested"]:
                    keep *= 1.0 - q.get(cr["fips"], 0.0)
            p = 1.0 - keep
        deaths[cr["fips"]] = deaths.get(cr["fips"], 0.0) + cr["deaths"] * p
    rates = {f: deaths.get(f, 0.0) / miles[f] * 1e8 for f in miles}
    return dict(catch=catch, missed=missed, untested=untested, deaths=deaths, rates=rates)


def backtest(obs):
    """The State Patrol posts test every surviving driver, so each post's arrests can be completed as if it were an
    arrest-only agency and checked against its own unarrested drivers at .08 or more. Each post is held out: the
    catch rates (and the one unarrested share) come from every other complete agency. Returns each post's actual
    count, the catch-rate completion and the pooled share's, and the two root mean squared errors."""
    rows = [(cr["fips"], d) for cr in obs for d in cr["drivers"]
            if d["complete"] and d["tested"] and not d["killed"]]
    out = {}
    for f in sorted({f for f, d in rows if d["agency"] == "SP"}):
        mine = [d for g, d in rows if g == f and d["agency"] == "SP"]
        rest = [d for g, d in rows if not (g == f and d["agency"] == "SP")]
        catch = {}
        for b, lo, hi in ((0, 0.080, 0.150), (1, 0.150, 9.9)):
            r = [d for d in rest if lo <= d["bac"] < hi]
            catch[b] = sum(1 for d in r if d["arrested"]) / float(len(r))
        un = [d for d in rest if not d["arrested"]]
        share = sum(1 for d in un if d["bac"] >= 0.080) / float(len(un))
        actual = sum(1 for d in mine if not d["arrested"] and d["bac"] >= 0.080)
        ipw = sum(1.0 / catch[0 if d["bac"] < 0.150 else 1] - 1.0 for d in mine if d["arrested"] and d["bac"] >= 0.080)
        pooled = share * sum(1 for d in mine if not d["arrested"])
        out[f] = (actual, ipw, pooled)
    rmse = lambda k: (sum((v[k] - v[0]) ** 2 for v in out.values()) / len(out)) ** 0.5
    return out, rmse(1), rmse(2)


# ------------------------------------------------------------------------------------------------- the laboratory
def lab(target):
    return read_csv(find(target, "stl_specimen_extract"))


def sites(target):
    return {r["site_code"]: r for r in read_csv(find(target, "stl_collection_sites"))}


def _local(utc):
    t = dt.datetime.strptime(utc, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=dt.timezone.utc)
    return t.astimezone(CENTRAL).replace(tzinfo=None)


def arrests_2025(rows, dr, county_of=None, year_of=None, merge_jail=True):
    """Each 2025 impaired-driving arrest with a blood specimen: one per submitting agency and case number, its county,
    and whether its first specimen's whole-blood ethanol was .15 or more."""
    kits = {}
    for r in rows:
        if r["case_type"] != "DUI":
            continue
        k = kits.setdefault(r["lab_number"], dict(agency=r["submitting_agency"], case=r["agency_case_number"],
                                                  utc=r["collected_utc"], site=r["collection_site"],
                                                  specimen=r["specimen"], eth=None, rows=[]))
        k["rows"].append(r)
        if r["analyte"] == "ETHANOL" and r["report_status"] in ("FINAL", "AMENDED") and r["result"]:
            v = float(r["result"])
            k["eth"] = v / 1.16 if r["specimen"] == "SER" else v
    by_case = {}
    for ln, k in kits.items():
        by_case.setdefault((k["agency"], k["case"]), []).append(k)
    if merge_jail:
        # one draw submitted twice, by the arresting department and by the county jail under the sheriff's booking
        # number: the same collection site and minute under two agencies is one arrest, the department's
        first_at = {}
        for key, ks in by_case.items():
            k0 = min(ks, key=lambda k: k["utc"])
            first_at.setdefault((k0["site"], k0["utc"]), []).append(key)
        for spot, keys in first_at.items():
            if len(keys) > 1:
                keep = [key for key in keys if not key[0].startswith("SO")] or keys[:1]
                for key in keys:
                    if key not in keep[:1]:
                        del by_case[key]
    out = []
    for (agency, case), ks in by_case.items():
        ks.sort(key=lambda k: k["utc"])
        first = ks[0]
        when = _local(first["utc"]) if year_of is None else year_of(first)
        if (when.year if hasattr(when, "year") else when) != 2025:
            continue
        if county_of:
            f = county_of(agency, case, first)
        elif agency == "SPHQ000":
            f = str(dr["posts"][int(case[1:3])]["county_fips"])
        else:
            f = str(dr["agencies"][agency]["county_fips"])
        out.append(dict(agency=agency, case=case, fips=f, high=first["eth"] is not None and first["eth"] >= 0.150))
    return out


def ask_a(rows, dr, **kw):
    arr = arrests_2025(rows, dr, **kw)
    out = {}
    for a in arr:
        n, h = out.get(a["fips"], (0, 0))
        out[a["fips"]] = (n + 1, h + (1 if a["high"] else 0))
    return {f: (n, 100.0 * h / n) for f, (n, h) in out.items()}


# ------------------------------------------------------------------------------------------------- crash records
def crs(target):
    reps = read_csv(find(target, "crs_reports_2025"))
    units = read_csv(find(target, "crs_units_2025"))
    u = {}
    for r in units:
        u.setdefault((r["report_id"], r["version"]), []).append(r)
    return reps, u


def ask_b(target, census_crashes, dedupe=True, link_census=True, latest=True, link_near=True):
    reps, units = crs(target)
    by = {}
    for r in reps:
        by.setdefault(r["report_id"], []).append(r)
    stand = []
    for rid, vs in by.items():
        acc = [v for v in vs if v["version_status"] == "ACCEPTED"]
        v = max(acc, key=lambda v: int(v["version"])) if latest else min(vs, key=lambda v: int(v["version"]))
        us = units.get((rid, v["version"]), [])
        v = dict(v)
        v["units"] = us
        v["drivers"] = tuple(sorted((x["driver_age"], x["driver_sex"]) for x in us if x["unit_type"] == "MOTOR VEHICLE"))
        stand.append(v)
    fatal = {(c["agency"], c["report"]) for c in census_crashes if c["year"] == 2025}
    # the census cites a death investigation number where a person died after the report: a crash is also the
    # census's when its date, county and road agree and its time is within 15 minutes
    near = {}
    for c in census_crashes:
        if c["year"] == 2025:
            near.setdefault((c["date"].isoformat(), c["fips"], c.get("road", "")), []).append(c["time"])

    def in_census(v):
        if (v["agency_id"], v["report_number"]) in fatal:
            return True
        if not link_near:
            return False
        for t in near.get((v["crash_date"], v["county_fips"], v["road_name"]), []):
            if not t:
                return True
            t0 = dt.datetime.strptime(t, "%H:%M")
            t1 = dt.datetime.strptime(v["crash_time"], "%H:%M")
            if abs((t1 - t0).total_seconds()) <= 15 * 60:
                return True
        return False
    # one crash, two agencies' reports: same date, county, road and drivers, times within 15 minutes
    groups = []
    if dedupe:
        key = {}
        for v in stand:
            key.setdefault((v["crash_date"], v["county_fips"], v["road_name"], v["drivers"]), []).append(v)
        for k, vs in key.items():
            vs.sort(key=lambda v: v["crash_time"])
            cur = [vs[0]]
            for v in vs[1:]:
                t0 = dt.datetime.strptime(cur[-1]["crash_time"], "%H:%M")
                t1 = dt.datetime.strptime(v["crash_time"], "%H:%M")
                if abs((t1 - t0).total_seconds()) <= 15 * 60 and v["agency_id"] != cur[-1]["agency_id"]:
                    cur.append(v)
                else:
                    groups.append(cur)
                    cur = [v]
            groups.append(cur)
    else:
        groups = [[v] for v in stand]
    out = {}
    for g in groups:
        if link_census and any(in_census(v) for v in g):
            continue
        ok = any(v["public_road"] == "Y" and v["severity"] in ("A", "B", "C") and
                 any(x["unit_type"] == "MOTOR VEHICLE" and x["driver_condition"] == "HAD BEEN DRINKING"
                     for x in v["units"]) for v in g)
        if ok:
            f = g[0]["county_fips"]
            out[f] = out.get(f, 0) + 1
    return out, groups


def golden_csv(gdir):
    return read_csv(find(gdir, "county_impaired_driving_rates"))
