"""Independent verifier: reads only the shipped bundle (and the golden memo), shares no code with
the generator, recomputes every graded figure and checks the memo's tables against it."""
import os
import sys

import docx
import openpyxl
import pandas as pd

TASK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def main(target=None, golden=None):
    t = target or os.path.join(TASK, "target")
    g = golden or os.path.join(TASK, "golden")
    facts = pd.read_parquet(os.path.join(t, "card_week_sales_2026w15_w34.parquet"))
    alloc = pd.read_csv(os.path.join(t, "st26_offer_allocation.csv"), dtype=str)[["card_no", "cell"]]
    cards = pd.read_csv(os.path.join(t, "member_cards.csv"), dtype=str)
    memb = pd.read_csv(os.path.join(t, "memberships.csv"), dtype=str)
    br = pd.read_csv(os.path.join(t, "branches_2026-01.csv"), dtype=str)
    ws = openpyxl.load_workbook(os.path.join(t, "programme_plan_2027.xlsx"), data_only=True)["Active cards"]
    plan = sum(r[3] for r in ws.iter_rows(min_row=5, values_only=True) if r[0] and r[0] != "Region code")

    base = [f"2026-W{w:02d}" for w in range(15, 23)]
    run = [f"2026-W{w:02d}" for w in range(29, 35)]
    facts["net"] = facts["gross_margin_chf"] - facts["reward_cost_chf"]
    wk = facts.pivot_table(index="card_no", columns="iso_week", values="net", aggfunc="sum")
    wk = wk.reindex(alloc["card_no"]).fillna(0.0)
    d = alloc.set_index("card_no").copy()
    d["chg"] = wk[run].sum(axis=1) / 6 - wk[base].sum(axis=1) / 8
    d = d.join(cards.set_index("card_no")["membership_no"])
    same = d.groupby("membership_no")["cell"].transform("nunique") == 1
    s = d[same]

    def lift(frame):
        m = frame.groupby("cell")["chg"].mean()
        return {x: m[x] - m["K"] for x in "ABCD"}

    chain = lift(s)
    naive = lift(d)
    value = {x: chain[x] * 52 * plan for x in "ABCD"}
    order = sorted(chain, key=chain.get, reverse=True)

    # regions of the 1 July 2026 network (branch register plus the network notice)
    reg = br.set_index("branch_no")["region_code"].to_dict()
    reg.update({"3107": "SGS", "3112": "SGS", "3115": "SGS"})
    succ = {"2209": "6101", "4105": "4101"}
    mreg = memb.set_index("membership_no")["home_branch"].map(lambda b: reg[succ.get(b, b)])
    s = s.assign(region=s["membership_no"].map(mreg))
    grid = {r: lift(s[s["region"] == r]) for r in ["SGS", "FUE", "RHT", "TGO", "TOG", "APZ"]}

    print("plan total", plan)
    print("chain", {x: round(v, 4) for x, v in chain.items()})
    print("naive card-level", {x: round(v, 4) for x, v in naive.items()}, "-> files", max(naive, key=naive.get))
    print("pick", order[0], "runner-up", order[1], "gap %.4f" % (chain[order[0]] - chain[order[1]]))
    print("2027 values (CHF k)", {x: round(v / 1000) for x, v in value.items()})
    for r in grid:
        print(r, {x: round(v, 2) for x, v in grid[r].items()})

    # the memo's tables against the recomputation
    ok = True
    names = {"Frischbonus": "A", "Dreifachpunkte": "B", "Monatsziel": "C", "Wochenauswahl": "D"}
    regions = {"St. Gallen Stadt": "SGS", "Fürstenland": "FUE", "Rheintal": "RHT", "Thurgau Ost": "TGO",
               "Toggenburg": "TOG", "Appenzellerland": "APZ"}
    doc = docx.Document(os.path.join(g, "programme_offer_2027.docx"))
    t1, t2 = doc.tables[1], doc.tables[2]
    for row in t1.rows[1:]:
        cells = [c.text for c in row.cells]
        x = names[cells[0].split(" ")[0]]
        ok &= cells[1] == f"{chain[x]:.2f}"
        ok &= cells[2].replace("'", "") == str(round(value[x] / 1000))
    hdr = [names[c.text] for c in t2.rows[0].cells[1:]]
    for row in t2.rows[1:]:
        cells = [c.text for c in row.cells]
        src = chain if cells[0] == "All regions" else grid[regions[cells[0]]]
        for x, v in zip(hdr, cells[1:]):
            ok &= v == f"{src[x]:.2f}"
    head = next(p.text for p in doc.paragraphs if p.text.startswith("Recommendation:"))
    ok &= "Frischbonus" in head and f"CHF {round(value['A'] / 1000)} thousand" in head
    print("memo tables and first line match the recomputation:", ok)
    return ok


if __name__ == "__main__":
    sys.exit(0 if main() else 1)
