"""task131 generator: the filed documents and workbooks. The trial plan, the gazetted residential tariffs, the
2026 demand forecast, Metering Services' file notes, the March survey's instruments, the trial register and the
2027-28 generation availability (a file the solution never reads)."""
from datetime import date

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import params as P
import knobs as K
import postprocess

ORG = P.ORG


def _styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13, spaceAfter=4)
    h = ParagraphStyle("h", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, spaceBefore=8, spaceAfter=3)
    t = ParagraphStyle("t", parent=ss["Title"], fontName="Helvetica-Bold", fontSize=14, alignment=0, spaceAfter=4)
    small = ParagraphStyle("s", parent=body, fontSize=8, leading=10.5, textColor=colors.HexColor("#444444"))
    return body, h, t, small


def _footer(text):
    def draw(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(20 * mm, 10 * mm, text)
        canvas.drawRightString(190 * mm, 10 * mm, "Page %d" % doc.page)
        canvas.restoreState()
    return draw


def _pdf(path, title, author, footer, blocks, stamp):
    body, h, t, small = _styles()
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                            bottomMargin=18 * mm, title=title, author=author, creator=author, invariant=1)
    flow = []
    for kind, val in blocks:
        if kind == "title":
            flow.append(Paragraph(val, t))
        elif kind == "h":
            flow.append(Paragraph(val, h))
        elif kind == "p":
            flow.append(Paragraph(val, body))
        elif kind == "small":
            flow.append(Paragraph(val, small))
        elif kind == "sp":
            flow.append(Spacer(1, val))
        elif kind == "table":
            rows, widths = val
            cell = ParagraphStyle("c", parent=body, fontName="Helvetica", fontSize=7.8, leading=9.6, spaceAfter=0)
            cellb = ParagraphStyle("cb", parent=cell, fontName="Helvetica-Bold")
            rows = [[Paragraph(str(x), cellb if r == 0 else cell) for x in row] for r, row in enumerate(rows)]
            tb = Table(rows, colWidths=widths, repeatRows=1)
            tb.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 7.8),
                                    ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 7.8),
                                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#9a9a9a")),
                                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6eaef")),
                                    ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            flow.append(tb)
    doc.build(flow, onFirstPage=_footer(footer), onLaterPages=_footer(footer))
    postprocess.scrub(path, ORG, stamp)


def _c(x):
    return "%.4f" % x


# ------------------------------------------------------------------------------------------------ trial plan
def trial_plan(path):
    ex = lambda v: _c(v / 1.1)
    pa, pb, pc = P.PRICES["A"], P.PRICES["B"], P.PRICES["C"]
    rows = [["Design", "Peak (weekdays)", "Shoulder (weekdays)", "Peak c/kWh", "Shoulder c/kWh", "Off-peak c/kWh", "Supply c/day"],
            ["A  Evening 3 to 9", "3 pm to 9 pm", "none", ex(pa["peak"]), "", ex(pa["off"]), ex(pa["supply"])],
            ["B  Evening 4 to 8", "4 pm to 8 pm", "none", ex(pb["peak"]), "", ex(pb["off"]), ex(pb["supply"])],
            ["C  Stepped", "4 pm to 8 pm", "2 pm to 4 pm and 8 pm to 10 pm", ex(pc["peak"]), ex(pc["shoulder"]), ex(pc["off"]), ex(pc["supply"])]]
    _pdf(path, "Summer pricing trial 2025-26: trial plan", "Load Research", "Murrandi Regional Power, load research, summer pricing trial 2025-26, plan", [
        ("title", "Summer pricing trial 2025-26: trial plan"),
        ("small", "Murrandi Regional Power, load research. Issued 3 November 2025, version 2 (version 1 of 18 September "
                  "2025 did not include Design C's shoulder prices). Owner: Peter Huang, load research lead."),
        ("h", "1. Why we are running it"),
        ("p", "The corporation advises the Minister each year on the residential tariffs for the following financial year. "
              "Before that advice for 2027-28 we want to see how Home Plan households use electricity when the price "
              "is higher in the late afternoon and evening, under three designs that differ in when the higher price "
              "applies and how sharply it steps."),
        ("h", "2. The designs"),
        ("table", (rows, [29 * mm, 22 * mm, 34 * mm, 17 * mm, 19 * mm, 19 * mm, 17 * mm])),
        ("small", "Prices exclude GST. Peak and shoulder times apply on weekdays other than public holidays; every other "
                  "time is off-peak. Times are Western Australian time. The Home Plan price and supply charge are those "
                  "gazetted for 2025-26."),
        ("h", "3. Who takes part"),
        ("p", "Home Plan households with an interval meter and at least twelve months of interval reads at 31 October "
              "2025, other than life-support registrations. Load research drew 1,600 households at random and placed "
              "them in blocks of four with similar weekday use between 4 pm and 9 pm in the 2024-25 summer and the same "
              "rooftop PV status. Within each block one household was assigned to each design and one to the control "
              "group, by lot. The register lists every household with its block and assignment."),
        ("h", "4. Dates"),
        ("p", "Letters to design households from 10 November 2025. Designs apply from 1 December 2025 to 28 February "
              "2026; every design household returns to the Home Plan on 1 March 2026. Control households stay on the "
              "Home Plan and are not contacted."),
        ("h", "5. Leaving a design"),
        ("p", "A household may return to the Home Plan at any time by calling us or through My Account; the change "
              "takes effect from the day after the request, and a request can be withdrawn within ten days. A household "
              "that registers for life support during the trial is returned to the Home Plan."),
        ("h", "6. How the trial will be read"),
        ("p", "Each design's households will be compared with the control households as assigned, whatever they did "
              "after assignment, because a default tariff would carry the same right to return to the Home Plan. Interval "
              "reads for every trial premise come from Metering Services' monthly NEM12 extracts; the billing system's "
              "account history records each account's tariff by register and date. Design households will be surveyed "
              "in March 2026."),
    ], "2025-11-03")


# ------------------------------------------------------------------------------------------------ tariffs
def tariffs(path):
    ph, pa, pb, pc, cl = P.PRICES["HOME"], P.PRICES["A"], P.PRICES["B"], P.PRICES["C"], P.PRICES["CL"]
    rows = [["Tariff", "Code", "Energy charge (c/kWh)", "Supply charge (c/day)"],
            ["Home Plan", P.CODE["HOME"], _c(ph["flat"]), _c(ph["supply"])],
            ["Daytime Saver", P.CODE["DS"], "11 am to 3 pm 9.1300; 3 pm to 9 pm 54.2671; other times 25.4108", "112.4381"],
            ["Hot Water Off-Peak", P.CODE["CL"], _c(cl["flat"]), "nil"],
            ["Trial Design A", P.CODE["A"], "peak %s; off-peak %s" % (_c(pa["peak"]), _c(pa["off"])), _c(pa["supply"])],
            ["Trial Design B", P.CODE["B"], "peak %s; off-peak %s" % (_c(pb["peak"]), _c(pb["off"])), _c(pb["supply"])],
            ["Trial Design C", P.CODE["C"], "peak %s; shoulder %s; off-peak %s" % (_c(pc["peak"]), _c(pc["shoulder"]), _c(pc["off"])), _c(pc["supply"])]]
    hol = [["Public holiday", "Date"]] + [[v, k.strftime("%A %d %B %Y").replace(" 0", " ")] for k, v in sorted(P.HOLIDAYS.items())]
    _pdf(path, "Residential charges 2025-26", "Murrandi Regional Power", "Murrandi Regional Power (Charges) By-laws 2025, Schedule 1, residential", [
        ("title", "Schedule 1: residential charges, 2025-26"),
        ("small", "Murrandi Regional Power (Charges) By-laws 2025, as amended by the Murrandi Regional Power (Charges) "
                  "Amendment By-laws (No. 2) 2025 gazetted 31 October 2025. Charges apply from 1 July 2025 unless stated."),
        ("h", "1. Charges"),
        ("table", (rows, [33 * mm, 14 * mm, 88 * mm, 33 * mm])),
        ("small", "All charges include GST. The trial designs apply from 1 December 2025 to 28 February 2026 to the "
                  "accounts the corporation places on them for its summer pricing trial."),
        ("h", "2. Times"),
        ("p", "Trial Design A peak time is 3 pm to 9 pm; Trial Design B and Trial Design C peak time is 4 pm to 8 pm; "
              "Trial Design C shoulder time is 2 pm to 4 pm and 8 pm to 10 pm. Peak, shoulder and Daytime Saver times "
              "apply on weekdays other than public holidays, and every other time is off-peak. Times are Western "
              "Australian Standard Time."),
        ("table", (hol, [60 * mm, 60 * mm])),
        ("h", "3. Hot Water Off-Peak"),
        ("p", "Hot Water Off-Peak applies to a separately metered hot water circuit that the corporation energises "
              "between 11 pm and 7 am."),
        ("h", "4. Renewable energy buyback"),
        ("p", "Electricity exported to the network from a customer's rooftop PV system is credited at %s cents per "
              "kWh (code BB1) whatever the customer's tariff." % _c(P.PRICES["BUYBACK"])),
    ], "2025-10-31")


# ------------------------------------------------------------------------------------------------ meter data notes
def meter_notes(path):
    _pdf(path, "Load research extracts: file notes", "Metering Services", "Murrandi Regional Power, Metering Services, load research extracts, file notes", [
        ("title", "Load research extracts: file notes"),
        ("small", "Murrandi Regional Power, Metering Services. Version 3, 12 February 2024. Owner: Nathaniel Maxwell."),
        ("p", "Load research receives a monthly extract of interval reads for the meters it names. Each extract is a "
              "NEM12 file in the AEMO meter data file format. Each extract holds the days of its month as first "
              "delivered and every day sent again since the previous extract."),
        ("h", "Record types"),
        ("table", ([["Record", "Holds"],
                    ["100", "file header: version, file date and time, sender, recipient"],
                    ["200", "one meter register: NMI, NMI configuration, register ID, NMI suffix, data stream, meter serial number, unit of measure, interval length in minutes"],
                    ["300", "one day of interval reads for the 200 record above it: interval date, the interval values, quality method, reason code, reason description, update date-time, MSATS load date-time"],
                    ["400", "quality by interval range, following a 300 record whose quality method is V"],
                    ["900", "end of file"]], [18 * mm, 150 * mm])),
        ("h", "Values"),
        ("p", "Interval values are kWh in the interval. A 30-minute meter has 48 values a day and a 15-minute meter 96. "
              "Interval 1 begins at midnight Western Australian Standard Time; Western Australia does not observe "
              "daylight saving. NMI suffix E1 is general supply import, E2 controlled-load import and B1 export to the "
              "network."),
        ("h", "Quality"),
        ("table", ([["Quality method", "Meaning"],
                    ["A", "actual reads"], ["E", "estimated"], ["S", "substituted"], ["F", "final substitute"],
                    ["N", "null: no reads for the interval"], ["V", "variable: the quality of each interval range is in the 400 records that follow"]],
                   [30 * mm, 138 * mm])),
        ("p", "The update date-time records when the reads were last updated; a day sent again carries a later update "
              "date-time."),
    ], "2024-02-12")


# ------------------------------------------------------------------------------------------------ survey instruments
def survey_instruments(path):
    web = [["Q", "Question", "Answers (code)"],
           ["Q1", "How well do you understand when your design's higher prices apply?", "Very well (1), Well (2), Somewhat (3), Not very well (4), Not at all (5)"],
           ["Q2", "Compared with the Home Plan, how do you think your bills changed on the design?", "Went down a lot (1), Went down a little (2), About the same (3), Went up a little (4), Went up a lot (5)"],
           ["Q3", "How comfortable was your home in the late afternoon and evening on the design?", "Very comfortable (1), Comfortable (2), Neither (3), Uncomfortable (4), Very uncomfortable (5)"],
           ["Q4", "If your design became your tariff from July 2027, how likely would you be to stay on it?", "Very likely (1), Likely (2), Unsure (3), Unlikely (4), Very unlikely (5)"],
           ["Q5", "How likely would you be to recommend the design to a friend?", "Very likely (1), Likely (2), Unsure (3), Unlikely (4), Very unlikely (5)"]]
    cati = [["Q", "Interviewer reads", "Code as read"],
            ["Q1", "How well do you understand when your design's higher prices apply? Would you say very well, well, somewhat, not very well or not at all?", "1 to 5 in the order read"],
            ["Q2", "Compared with the Home Plan, do you think your bills went down a lot, went down a little, stayed about the same, went up a little or went up a lot?", "1 to 5 in the order read"],
            ["Q3", "How comfortable was your home in the late afternoon and evening: very comfortable, comfortable, neither, uncomfortable or very uncomfortable?", "1 to 5 in the order read"],
            ["Q4", "If your design became your tariff from July 2027, would you be very unlikely, unlikely, unsure, likely or very likely to stay on it?", "1 to 5 in the order read"],
            ["Q5", "How likely would you be to recommend the design to a friend: very likely, likely, unsure, unlikely or very unlikely?", "1 to 5 in the order read"]]
    _pdf(path, "Summer pricing trial: March 2026 survey instruments", "Load Research", "Murrandi Regional Power, summer pricing trial, March 2026 survey", [
        ("title", "Summer pricing trial: March 2026 survey"),
        ("small", "Murrandi Regional Power, load research, with Customer Insights. Fieldwork 2 to 24 March 2026."),
        ("h", "How the survey was run"),
        ("p", "An email invitation went on 2 March 2026 to every design household, with a reminder on 9 March. From 12 to "
              "24 March Customer Insights telephoned households that had not answered online. A response counts once it "
              "is submitted; where a household submitted more than once, its first submitted response counts. The "
              "extract lists every response started, submitted or not, by trial ID."),
        ("h", "Online form"),
        ("table", (web, [10 * mm, 82 * mm, 76 * mm])),
        ("h", "Telephone script"),
        ("table", (cati, [10 * mm, 118 * mm, 40 * mm])),
    ], "2026-03-01")


# ------------------------------------------------------------------------------------------------ workbooks
def _hdr(ws, cols, row=1):
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE4EE")
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _widths(ws, widths):
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w


def _save(wb, path, who, created, modified):
    wb.save(path)
    postprocess.scrub(path, who, modified[:10])
    postprocess.finish_office(path, created, modified, ORG)


def curve10():
    """the 10% POE design day: warmer through the night as well, and much hotter through the afternoon and evening"""
    c = np.array(K.CURVE)
    return list(np.round(c * 1.03 + np.clip(c - 140.0, 0, None) * 0.18, 1))


def design_temps(tmax, tmin):
    import weather as WX
    return WX.trace(tmax, tmin, tmax - 0.4)


def forecast(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cover"
    for r, (a, b) in enumerate([
        ("Murrandi system demand forecast 2026", ""),
        ("Issued", "14 August 2026"), ("Version", "1.2"), ("Owner", "Craig Jackson, system planning"),
        ("Supersedes", "Murrandi system demand forecast 2025 (version 1.1)"),
        ("Basis", "Customers on the tariffs they hold in 2026; sent-out demand at the power station; reference weather at Murrandi Aero"),
    ], start=1):
        ws.cell(row=r, column=1, value=a).font = Font(bold=(r == 1))
        ws.cell(row=r, column=2, value=b)
    _widths(ws, [26, 90])

    ws = wb.create_sheet("Summary")
    _hdr(ws, ["Summer", "Maximum demand 50% POE (MW)", "Maximum demand 10% POE (MW)", "Energy sent out (GWh)"])
    c10 = max(curve10())
    md50 = {"2026-27": 211.3, "2027-28": P.MD50, "2028-29": 217.2, "2029-30": 219.9, "2030-31": 222.1}
    md10 = {"2026-27": 230.6, "2027-28": float(c10), "2028-29": 237.3, "2029-30": 240.0, "2030-31": 242.8}
    energy = {"2026-27": 1018.4, "2027-28": 1031.3, "2028-29": 1041.7, "2029-30": 1055.2, "2030-31": 1064.6}
    for r, (y, v) in enumerate(md50.items(), start=2):
        ws.cell(row=r, column=1, value=y)
        ws.cell(row=r, column=2, value=v)
        ws.cell(row=r, column=3, value=md10[y])
        ws.cell(row=r, column=4, value=energy[y])
    _widths(ws, [12, 26, 26, 22])

    ws = wb.create_sheet("Design day 2027-28")
    _hdr(ws, ["Trading interval", "Interval start", "50% POE demand (MW)", "50% POE temperature (C)", "10% POE demand (MW)", "10% POE temperature (C)"])
    t50 = design_temps(P.DESIGN_TMAX, P.DESIGN_TMIN)
    t10 = design_temps(46.3, 31.4)
    c10l = curve10()
    for k in range(48):
        ws.cell(row=k + 2, column=1, value=k + 1)
        ws.cell(row=k + 2, column=2, value="%02d:%02d" % (k // 2, 30 * (k % 2)))
        ws.cell(row=k + 2, column=3, value=K.CURVE[k])
        ws.cell(row=k + 2, column=4, value=round(float(t50[k]), 1))
        ws.cell(row=k + 2, column=5, value=c10l[k])
        ws.cell(row=k + 2, column=6, value=round(float(t10[k]), 1))
    _widths(ws, [16, 14, 20, 22, 20, 22])
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Customers")
    _hdr(ws, ["At 1 July", "Home Plan", "Daytime Saver", "Residential total", "Non-residential"])
    data = [("2026", 41592, 2716), ("2027", P.N_FLAT_2728, P.N_DS_2728), ("2028", 42071, 3168), ("2029", 42255, 3391),
            ("2030", 42418, 3605)]
    for r, (y, a, b) in enumerate(data, start=2):
        ws.cell(row=r, column=1, value=y)
        ws.cell(row=r, column=2, value=a)
        ws.cell(row=r, column=3, value=b)
        ws.cell(row=r, column=4, value=a + b)
        ws.cell(row=r, column=5, value=[4120, 4143, 4151, 4176, 4188][r - 2])
    _widths(ws, [10, 12, 14, 18, 16])

    ws = wb.create_sheet("Method")
    notes = [
        "Definitions",
        "The summer maximum demand at 50% probability of exceedance (50% POE) is the highest half-hourly sent-out "
        "demand on the 50% POE design day; it has an even chance of being exceeded in the summer. The 10% POE maximum "
        "demand is the same on the 10% POE design day.",
        "Demand is MW sent out from the Murrandi power station, averaged over the half-hour, and includes the network's "
        "losses. Rooftop PV generation is netted from customers' demand.",
        "Trading interval 1 is 00:00 to 00:30 Western Australian Standard Time; trading interval 48 ends at midnight.",
        "Each design day is a weekday. Its temperatures are the forecast's reference weather at Murrandi Aero: a daily "
        "maximum of 44.6 C for the 50% POE day and 46.3 C for the 10% POE day.",
        "Customer numbers are those forecast on each tariff at 1 July of the year. Every Home Plan customer has had an "
        "interval meter since the 2024 meter replacement programme.",
        "The forecast assumes customers stay on the tariffs they hold in 2026. A change in customers' demand is carried "
        "into the forecast as metered, without a loss adjustment.",
    ]
    for r, t in enumerate(notes, start=1):
        c = ws.cell(row=r, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if r == 1:
            c.font = Font(bold=True)
    _widths(ws, [120])
    _save(wb, path, "Craig Jackson", "2026-07-28T03:12:00Z", "2026-08-14T06:40:00Z")


def register_book(df, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Register"
    cols = list(df.columns)
    _hdr(ws, cols)
    for r, row in enumerate(df.itertuples(index=False), start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(row=r, column=j, value=(v.item() if hasattr(v, "item") else v))
    _widths(ws, [13, 12, 13, 13, 8, 11, 11, 17, 24])
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("Notes")
    for r, t in enumerate([
        "Summer pricing trial 2025-26: register of trial households",
        "One row per trial household, the customer named on the account at enrolment on 31 October 2025. arm is the household's assignment by lot within its block.",
        "summer_2024_25_4to9pm_kw is the household's mean weekday demand between 4 pm and 9 pm, December 2024 to February 2025, used to form the blocks.",
        "rooftop_pv and hot_water_circuit are as recorded at enrolment.",
    ], start=1):
        ws2.cell(row=r, column=1, value=t).font = Font(bold=(r == 1))
    _widths(ws2, [130])
    _save(wb, path, "Peter Huang", "2025-10-31T01:05:00Z", "2025-11-07T05:22:00Z")


def generation(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Units 2027-28"
    _hdr(ws, ["Unit", "Type", "Nameplate (MW)", "Summer firm (MW)", "Planned outage 2027-28", "Notes"])
    units = [("MPS G1", "gas engine", 18.4, 16.9, "none", ""), ("MPS G2", "gas engine", 18.4, 16.9, "May 2028", ""),
             ("MPS G3", "gas engine", 18.4, 16.6, "none", ""), ("MPS G4", "gas engine", 18.4, 16.6, "none", ""),
             ("MPS G5", "gas engine", 18.4, 16.4, "April 2028", ""), ("MPS G6", "gas engine", 18.4, 16.4, "none", ""),
             ("MPS G7", "gas engine", 9.7, 8.8, "none", ""), ("MPS G8", "gas engine", 9.7, 8.8, "none", ""),
             ("MPS GT1", "open-cycle gas turbine", 31.0, 24.8, "none", "derates above 40 C"),
             ("MPS GT2", "open-cycle gas turbine", 31.0, 24.8, "June 2028", "derates above 40 C"),
             ("Murrandi Solar Farm", "single-axis solar", 40.0, 6.0, "none", "firm value at 6 pm"),
             ("MPS BESS", "battery", 20.0, 18.5, "none", "two-hour rating")]
    for r, u in enumerate(units, start=2):
        for j, v in enumerate(u, start=1):
            ws.cell(row=r, column=j, value=v)
    _widths(ws, [14, 22, 15, 16, 22, 22])
    _save(wb, path, "Craig Jackson", "2026-08-02T02:10:00Z", "2026-08-21T04:31:00Z")
