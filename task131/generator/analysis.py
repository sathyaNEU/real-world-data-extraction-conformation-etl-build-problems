"""task131 generator: the golden analysis, run on the shipped files only.

Parses the four NEM12 extracts (latest version of each meter-day, an exchange day summed across the two meters,
15-minute reads paired into half-hours), follows each trial household by its customer number through the account
history, and computes the intention-to-treat difference against the control households by half-hour; then each
design's change to the 2027-28 50% POE maximum demand, and the three asks. Also computes the wrong rungs, so the
checks can assert them on the same parsed data."""
import os
from collections import defaultdict
from datetime import date, datetime, timedelta

import numpy as np
import pandas as pd
from openpyxl import load_workbook

import params as P

NEMS = ["MRP_NEM12_202511.csv", "MRP_NEM12_202512.csv", "MRP_NEM12_202601.csv", "MRP_NEM12_202602.csv"]
D0 = P.DATA_FROM


def parse_nem(tgt, files=NEMS, dedup="latest_per_serial"):
    """(nmi, register) -> (NDAY, 48) kWh per half-hour; versions resolved per meter serial and summed across serials.
    dedup='nmi_reg_date' keeps one record per NMI, register and day (the latest), whatever the serial (the reflex);
    dedup='none' sums every record."""
    recs = defaultdict(dict)          # (nmi, reg, serial, day) -> (update, values)
    allrecs = defaultdict(list)
    for fn in files:
        nmi = reg = serial = None
        ilen = 30
        with open(os.path.join(tgt, fn)) as fh:
            for line in fh:
                if line.startswith("300"):
                    f = line.rstrip("\r\n").split(",")
                    nint = 96 if ilen == 15 else 48
                    vals = np.array(f[2:2 + nint], dtype=float)
                    if nint == 96:
                        vals = vals[0::2] + vals[1::2]
                    upd = f[2 + nint + 4]
                    day = (date(int(f[1][:4]), int(f[1][4:6]), int(f[1][6:8])) - D0).days
                    k = (nmi, reg, serial, day)
                    if dedup == "none":
                        allrecs[(nmi, reg, day)].append(vals)
                    else:
                        old = recs[(nmi, reg)].get((serial, day))
                        if old is None or upd > old[0]:
                            recs[(nmi, reg)][(serial, day)] = (upd, vals)
                elif line.startswith("200"):
                    f = line.rstrip("\r\n").split(",")
                    nmi, reg, serial, ilen = f[1], f[4], f[6], int(f[8])
    out = {}
    if dedup == "none":
        for (nmi, reg, day), vs in allrecs.items():
            arr = out.setdefault((nmi, reg), np.full((P.NDAY, 48), np.nan))
            arr[day] = np.sum(vs, axis=0)
        return out
    for key, d in recs.items():
        arr = np.full((P.NDAY, 48), np.nan)
        if dedup == "nmi_reg_date":
            best = {}
            for (serial, day), (upd, vals) in d.items():
                if day not in best or upd > best[day][0]:
                    best[day] = (upd, vals)
            for day, (upd, vals) in best.items():
                arr[day] = vals
        else:
            for (serial, day), (upd, vals) in d.items():
                arr[day] = vals if np.isnan(arr[day]).all() else arr[day] + vals
        out[key] = arr
    return out


def load_history(tgt):
    h = pd.read_csv(os.path.join(tgt, "account_tariff_history_2026-03-12.csv"), dtype=str, keep_default_na=False)
    for c in ("valid_from", "valid_to", "keyed_on"):
        h[c] = [date.fromisoformat(x) if x else None for x in h[c]]
    return h


def load_register(tgt):
    wb = load_workbook(os.path.join(tgt, "summer_pricing_trial_register.xlsx"), read_only=True)
    ws = wb["Register"]
    rows = list(ws.iter_rows(values_only=True))
    df = pd.DataFrame(rows[1:], columns=rows[0])
    for c in ("trial_id", "customer_id", "account_id", "nmi", "arm"):
        df[c] = df[c].astype(str)
    return df


def load_forecast(tgt):
    wb = load_workbook(os.path.join(tgt, "murrandi_demand_forecast_2026.xlsx"), read_only=True)
    ws = wb["Design day 2027-28"]
    rows = list(ws.iter_rows(values_only=True))[1:]
    c50 = np.array([r[2] for r in rows], dtype=float)
    t50 = np.array([r[3] for r in rows], dtype=float)
    c10 = np.array([r[4] for r in rows], dtype=float)
    ws = wb["Customers"]
    cust = {str(r[0]): r for r in list(ws.iter_rows(values_only=True))[1:]}
    return c50, t50, c10, int(cust["2027"][1]), int(cust["2027"][3])


def weather(tgt):
    w = pd.read_csv(os.path.join(tgt, "IDCJAC0010_013027_1800_Data.csv"), dtype=str, keep_default_na=False)
    out = {}
    for _, r in w.iterrows():
        if r["Maximum temperature (Degree C)"]:
            out[date(int(r["Year"]), int(r["Month"]), int(r["Day"]))] = float(r["Maximum temperature (Degree C)"])
    return out


def premises(hist, reg):
    """for each trial household (by customer number) the premise it occupied on each day of the data window:
    list of (nmi, from_day_index, to_day_index inclusive) on accounts held by that customer"""
    e1 = hist[hist.register == "E1"]
    by_cust = defaultdict(list)
    for _, r in e1.iterrows():
        by_cust[r.customer_id].append(r)
    out = {}
    for _, t in reg.iterrows():
        spans = []
        for r in by_cust[t.customer_id]:
            f = r.valid_from
            to = r.valid_to
            if to is not None and to < f:
                continue                  # a cancelled change
            spans.append((r.nmi, f, to))
        # merge to premise spans
        pm = defaultdict(lambda: [None, None])
        for nmi, f, to in spans:
            a, b = pm[nmi]
            pm[nmi] = [f if a is None or f < a else a, None if (to is None or b is None and a is not None and b is None) else max(to, b) if b else to]
        res = []
        for nmi, (f, to) in pm.items():
            lo = max((f - D0).days, 0)
            hi = P.NDAY - 1 if to is None else min((to - D0).days, P.NDAY - 1)
            if hi >= lo:
                res.append((nmi, lo, hi))
        out[t.trial_id] = sorted(res, key=lambda x: x[1])
    return out


def household_panel(reads, prem, reg):
    """(n, NDAY, 48) net demand in average kW, NaN where the household was not at a premise"""
    tids = list(reg.trial_id)
    X = np.full((len(tids), P.NDAY, 48), np.nan)
    for k, tid in enumerate(tids):
        for nmi, lo, hi in prem[tid]:
            e1 = reads.get((nmi, "E1"))
            if e1 is None:
                continue
            v = e1[lo:hi + 1].copy()
            e2 = reads.get((nmi, "E2"))
            b1 = reads.get((nmi, "B1"))
            if e2 is not None:
                v = v + np.nan_to_num(e2[lo:hi + 1])
            if b1 is not None:
                v = v - np.nan_to_num(b1[lo:hi + 1])
            X[k, lo:hi + 1] = v * 2.0
    return tids, X


def itt(X, arms, days, a):
    ia = np.where(arms == a)[0]
    ic = np.where(arms == "CONTROL")[0]
    out = []
    for j in days:
        out.append(np.nanmean(X[ia, j], axis=0) - np.nanmean(X[ic, j], axis=0))
    return np.mean(out, axis=0)


def md_change(curve, n, eff):
    new = curve + n * eff / 1000.0
    o = np.argsort(-new)
    return float(curve.max() - new[o[0]]), int(o[0]), float(new[o[0]] - new[o[1]])


def run(tgt, reads=None):
    reg = load_register(tgt)
    hist = load_history(tgt)
    c50, t50, c10, n_flat, n_res = load_forecast(tgt)
    wx = weather(tgt)
    reads = reads or parse_nem(tgt)
    prem = premises(hist, reg)
    tids, X = household_panel(reads, prem, reg)
    arms = np.array(reg.arm)
    work = [j for j, d in enumerate(P.DAYS) if P.trial_day(d) and P.workday(d)]
    tmax_design = float(max(t50))
    hot = [j for j in work if wx.get(P.DAYS[j], 0) >= 43.0]
    R = {"hot_days": [P.DAYS[j].isoformat() for j in hot], "n_flat": n_flat, "md50": float(c50.max()),
         "md50_ti": int(np.argmax(c50)) + 1, "design_tmax": tmax_design}
    win = {"A": (30, 42), "B": (32, 40), "C": (32, 40)}
    for a in P.DESIGNS:
        lab = "DESIGN " + a
        eh = itt(X, arms, hot, lab)
        ea = itt(X, arms, work, lab)
        ch, ti, gap = md_change(c50, n_flat, eh)
        R[a] = {"eff_hot": eh, "eff_avg": ea, "change": ch, "md_ti": ti + 1, "md_gap": gap, "md_mw": float(c50.max() - ch),
                "R0": float(-ea[win[a][0]:win[a][1]].mean() * n_flat / 1000), "R0b": float(-ea[36] * n_flat / 1000),
                "R1": float(-eh[36] * n_flat / 1000), "R2": md_change(c50, n_flat, ea)[0], "R2_ti": md_change(c50, n_flat, ea)[1] + 1}
    R["X"], R["arms"], R["tids"], R["reg"], R["hist"], R["prem"], R["reads"], R["c50"], R["c10"] = X, arms, tids, reg, hist, prem, reads, c50, c10
    R["wx"], R["work"], R["hot"] = wx, work, hot
    order = sorted(P.DESIGNS, key=lambda a: -R[a]["change"])
    R["call"] = order[0] if R[order[0]]["change"] > 0 else "none"
    R["order"] = order
    return R


# ------------------------------------------------------------------------------------------------ the asks
def ask_retention(hist, reg, on=date(2026, 2, 28)):
    """per design: still billed on it, back on the Home Plan at own request, no longer our customers (customer chain)"""
    e1 = hist[hist.register == "E1"]
    out = {}
    for a in P.DESIGNS:
        code = P.CODE[a]
        still = back = gone = 0
        for _, t in reg[reg.arm == "DESIGN " + a].iterrows():
            rows = e1[e1.customer_id == t.customer_id]
            live = [r for _, r in rows.iterrows() if r.valid_from <= on and (r.valid_to is None or r.valid_to >= on)
                    and not (r.valid_to is not None and r.valid_to < r.valid_from)]
            if not live:
                gone += 1
                continue
            r = live[0]
            if r.tariff_code == code:
                still += 1
                continue
            if r.tariff_code == P.CODE["HOME"]:
                # how did the household leave its design? the change off the design on any of its accounts
                req = False
                rs = rows[(rows.tariff_code == P.CODE["HOME"])]
                for _, x in rs.iterrows():
                    if x.valid_to is not None and x.valid_to < x.valid_from:
                        continue
                    if x.start_reason in ("CUSTOMER REQUEST", "SELF SERVICE"):
                        req = True
                if req:
                    back += 1
        out[a] = {"still": still, "back": back, "gone": gone}
    return out


def charges_day(kwh, design, d):
    """GST-inclusive energy charges in cents for one day's 48 half-hourly kWh on a design, and on the Home Plan"""
    pr = P.PRICES
    home = float(np.sum(kwh) * pr["HOME"]["flat"])
    if not P.workday(d):
        des = float(np.sum(kwh) * pr[design]["off"])
        return des, home
    rate = np.full(48, pr[design]["off"])
    lo, hi = P.WIN[design]["peak"]
    rate[lo:hi] = pr[design]["peak"]
    for a, b in P.WIN[design].get("shoulder", []):
        rate[a:b] = pr[design]["shoulder"]
    return float(np.sum(kwh * rate)), home


def ask_bills(hist, reg, reads):
    """per design: households charged more over the trial than the Home Plan for the same general-supply reads, and
    the mean difference per household in dollars"""
    e1 = hist[hist.register == "E1"]
    out = {}
    per = {}
    for a in P.DESIGNS:
        code = P.CODE[a]
        diffs = []
        for _, t in reg[reg.arm == "DESIGN " + a].iterrows():
            rows = e1[(e1.customer_id == t.customer_id) & (e1.tariff_code == code)]
            tot = 0.0
            for _, r in rows.iterrows():
                if r.valid_to is not None and r.valid_to < r.valid_from:
                    continue
                f = max(r.valid_from, P.START)
                to = min(r.valid_to or P.END, P.END)
                arr = reads[(r.nmi, "E1")]
                d = f
                while d <= to:
                    j = (d - D0).days
                    des, home = charges_day(arr[j], a, d)
                    tot += des - home + (P.PRICES[a]["supply"] - P.PRICES["HOME"]["supply"])
                    d += timedelta(1)
            diffs.append(tot / 100.0)
            per[t.trial_id] = tot / 100.0
        diffs = np.array(diffs)
        out[a] = {"more": int(np.sum(diffs > 0.0)), "mean": float(np.mean(diffs))}
    return out, per


def ask_survey(tgt, reg):
    sv = pd.read_csv(os.path.join(tgt, "trial_survey_responses_2026-03.csv"), dtype=str, keep_default_na=False)
    sv = sv[sv.complete == "Y"].copy()
    sv = sv.sort_values(["submitted_at", "response_id"], kind="mergesort")
    first = sv.drop_duplicates("trial_id", keep="first")
    arm = dict(zip(reg.trial_id, reg.arm))
    out = {}
    for a in P.DESIGNS:
        sub = first[[arm.get(t) == "DESIGN " + a for t in first.trial_id]]
        sub = sub[sub.q4_stay != ""]
        q = sub.q4_stay.astype(int)
        likely = np.where(sub["mode"] == "WEB", q <= 2, q >= 4)
        out[a] = {"answered": int(len(sub)), "likely": int(likely.sum()), "share": float(100.0 * likely.mean())}
    return out
