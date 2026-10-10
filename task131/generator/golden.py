"""The golden reading, computed from the shipped files only, and the two deliverables."""
import datetime as dt
import json
import os
import sys

import numpy as np
import pandas as pd
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import world as W
import estimators as E

TASK = os.path.dirname(HERE)
OFFER_LABEL = {
    "A": "Frischbonus (fresh bonus)",
    "B": "Dreifachpunkte (triple points)",
    "C": "Monatsziel (monthly target)",
    "D": "Wochenauswahl (weekly picks)",
}
SHORT = {"A": "Frischbonus", "B": "Dreifachpunkte", "C": "Monatsziel", "D": "Wochenauswahl"}


def load(target):
    f = pd.read_parquet(os.path.join(target, "card_week_sales_2026w15_w34.parquet"))
    a = pd.read_csv(os.path.join(target, "st26_offer_allocation.csv"), dtype={"card_no": str})
    c = pd.read_csv(os.path.join(target, "member_cards.csv"), dtype={"card_no": str, "membership_no": str})
    m = pd.read_csv(os.path.join(target, "memberships.csv"), dtype={"membership_no": str, "home_branch": str})
    b = pd.read_csv(os.path.join(target, "branches_2026-01.csv"), dtype={"branch_no": str})
    wb = openpyxl.load_workbook(os.path.join(target, "programme_plan_2027.xlsx"), data_only=True)
    ws = wb["Active cards"]
    planned = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] in dict(W.REGIONS):
            planned[row[0]] = row[3]
    return f, a[["card_no", "cell"]], c, m, b, planned


def region_2027(m, b):
    """Home branch -> successor for the two branches closed on 30 June -> region in the 1 July network."""
    reg = b.set_index("branch_no")["region_code"].to_dict()
    reg.update({"3107": "SGS", "3112": "SGS", "3115": "SGS"})           # notice: Rorschach area to SGS
    succ = {"2209": "6101", "4105": "4101"}                              # notice: closures
    hb = m["home_branch"].map(lambda x: succ.get(x, x))
    return pd.Series(hb.map(reg).to_numpy(), index=m["membership_no"])


def figures(target):
    f, a, c, m, b, planned = load(target)
    chain = E.answer_net_same_cell(f, a, c)
    total_cards = sum(planned.values())
    value = {x: E.programme_value_2027(chain[x], total_cards) for x in W.OFFERS}
    pick = max(chain, key=chain.get)
    runner = sorted(chain, key=chain.get)[-2]
    grid = E.grid(f, a, c, region_2027(m, b))
    return {
        "chain": chain, "value_chf": value, "pick": pick, "runner_up": runner,
        "gap": chain[pick] - chain[runner], "grid": grid, "planned_total": total_cards,
    }


def write_docx(fig, path):
    import docx
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2.3)
        s.top_margin = Cm(1.5)
        s.bottom_margin = Cm(1.6)
    cp = d.core_properties
    cp.author = "Patrice Keller"
    cp.last_modified_by = "Patrice Keller"
    cp.title = "Member programme offer 2027"
    cp.created = cp.modified = dt.datetime(2026, 10, 12, 8, 15)
    cp.revision = 3
    for k in ("comments", "keywords", "subject", "category"):
        setattr(cp, k, "")

    def para(text, bold=False, size=None, after=6, italic=False):
        p = d.add_paragraph()
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(after)
        return p

    def shade(cell, hexcol):
        tcPr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:color"), "auto")
        sh.set(qn("w:fill"), hexcol)
        tcPr.append(sh)

    def table(rows, widths, bold_first=True, bold_rows=()):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                cell = t.cell(i, j)
                cell.text = str(v)
                cell.width = Cm(widths[j])
                for p in cell.paragraphs:
                    if j > 0:
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    for run in p.runs:
                        run.font.size = Pt(9.5)
                        run.bold = (i == 0 and bold_first) or i in bold_rows
                if i == 0:
                    shade(cell, "E4EAF1")
        return t

    ch, val, pick, run = fig["chain"], fig["value_chf"], fig["pick"], fig["runner_up"]
    k = lambda v: f"{round(v / 1000):,}".replace(",", "'")

    # page number in the footer
    fp = d.sections[0].footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run("Mitgliederprogramm · Budget 2027 · page ")
    r.font.size = Pt(8)
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        frun = fp.add_run()
        frun.font.size = Pt(8)
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        frun._r.append(el)

    para("Talwerk Genossenschaft · Mitgliederprogramm", italic=True, size=8.5, after=6)
    hdr_rows = [("To", "Geschäftsleitung (executive board), budget round 2027"),
                ("From", "Patrice Keller, member programme"),
                ("Date", "12 October 2026"),
                ("Subject", "Member offer line, budget 2027 (agenda item 6, 4 November)")]
    t = d.add_table(rows=4, cols=2)
    t.autofit = False
    t.columns[0].width, t.columns[1].width = Cm(2.4), Cm(13.6)
    for i, (a_, b_) in enumerate(hdr_rows):
        t.cell(i, 0).text, t.cell(i, 1).text = a_, b_
        for j in (0, 1):
            t.cell(i, j).width = Cm(2.4 if j == 0 else 13.6)
            for p_ in t.cell(i, j).paragraphs:
                for hrun in p_.runs:
                    hrun.font.size = Pt(9.5)
                    hrun.bold = (j == 0)
    para("", after=4)
    para(f"The {SHORT[pick]} goes on every card in 2027", bold=True, size=15, after=8)

    para(f"Recommendation: run the {OFFER_LABEL[pick]} on every member card in 2027. "
         f"Its 2027 programme value is CHF {k(val[pick])} thousand.", bold=True, size=11.5, after=10)

    para(f"The {SHORT[pick]} adds CHF {ch[pick]:.2f} of net margin per active card-week at run rate. "
         f"The {SHORT[run]} comes second at CHF {ch[run]:.2f}, so the {SHORT[pick]} is ahead by "
         f"CHF {fig['gap']:.2f} per card-week. The table puts every offer on the same footing, and the "
         f"2027 values are on the {fig['planned_total']:,} active cards in finance's plan.".replace(
             f"{fig['planned_total']:,}", f"{fig['planned_total']:,}".replace(",", "'")),
         after=8)

    rows = [("Offer", "Net margin per active card-week (CHF)", "2027 programme value (CHF thousands)")]
    for x in sorted(W.OFFERS, key=lambda x: -ch[x]):
        rows.append((OFFER_LABEL[x], f"{ch[x]:.2f}", k(val[x])))
    table(rows, (6.2, 5.0, 5.0), bold_rows=(1,))
    para("Source: card-week sales facts, 6 April to 23 August 2026; ST26 allocation register; programme plan "
         "2027 (version 2). Measured as in MP-3, revision 3.¹", size=8, italic=True, after=10)

    para("How we read the test", bold=True, after=2)
    para("The offer engine allocated cards, not memberships, and a membership's cards shop from one household "
         "budget. Where a membership's cards drew different cells, the card carrying an offer pulled baskets off "
         "its sibling cards. Read card by card, that moved spend shows up twice, once as a gain on the offer card "
         "and once as a loss on the card it came from, and none of it exists when every card carries the offer. "
         "So the figures above compare memberships whose cards all drew the same cell, which is what a full "
         "rollout looks like. Read card by card, the Wochenauswahl would have come out first.", after=6)
    para("The CRM team expected the Monatsziel to win, and it did move more spend than any other offer during "
         "the test. Much of that spend was moved onto the voucher card to reach the CHF 250 monthly target, and "
         "the vouchers took most of the margin it added. The Dreifachpunkte started strongly and fell back "
         "after the first six weeks, which is why the run-rate weeks decide it.", after=10)

    para("By region", bold=True, after=2)
    para("The regional managers set their 2027 targets from this table. Net margin per active card-week in CHF, "
         "run rate.²", size=9.5, after=4)
    hdr = ["Region"] + [SHORT[x] for x in W.OFFERS]
    rows = [hdr]
    for code, name in W.REGIONS:
        rows.append([name] + [f"{fig['grid'][code][x]:.2f}" for x in W.OFFERS])
    rows.append(["All regions"] + [f"{ch[x]:.2f}" for x in W.OFFERS])
    table(rows, (4.2, 3.0, 3.0, 3.0, 3.0), bold_rows=(len(rows) - 1,))
    para("Source: as above, by the membership's home branch.", size=8, italic=True, after=10)

    para("¹ Net margin is gross margin less reward cost. Baseline weeks 2026-W15 to W22, run-rate weeks W29 to W34. "
         "Every allocated card is an active card, and a week without a purchase counts as zero. Annual value is "
         "the per card-week figure times 52 times the plan's active cards, rounded once.", size=8, after=2)
    para("² Regions of the network from 1 July 2026: Rorschach Hafen, Goldach and Thal in St. Gallen Stadt; "
         "Degersheim members with Herisau in Appenzellerland; Egnach members with Arbon.", size=8, after=2)
    d.save(path)


def write_png(fig, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.colors import LinearSegmentedColormap
    from matplotlib.patches import Rectangle

    ramp = ["#cde2fb", "#9ec5f4", "#6da7ec", "#3987e5", "#256abf", "#184f95"]
    cmap = LinearSegmentedColormap.from_list("blue", ramp)
    regions = [n for _, n in W.REGIONS]
    codes = [c for c, _ in W.REGIONS]
    rows = ["All regions"] + regions
    vals = np.array([[fig["chain"][x] for x in W.OFFERS]] +
                    [[fig["grid"][c][x] for x in W.OFFERS] for c in codes])
    pick, run = fig["pick"], fig["runner_up"]

    plt.rcParams["font.family"] = "DejaVu Sans"
    figr, ax = plt.subplots(figsize=(10, 6.4), dpi=150)
    vmin, vmax = 0.0, 0.6
    ax.imshow(vals, cmap=cmap, vmin=vmin, vmax=vmax, aspect="auto")
    for i in range(vals.shape[0]):
        for j in range(vals.shape[1]):
            v = vals[i, j]
            ink = "#ffffff" if (v - vmin) / (vmax - vmin) > 0.62 else "#1d1d1b"
            ax.text(j, i, f"{v:.2f}", ha="center", va="center", fontsize=11.5, color=ink,
                    fontweight="bold" if i == 0 else "normal")
    ax.set_xticks(range(4))
    ax.set_xticklabels([f"{SHORT[x]}" + ("\n(runner-up)" if x == run else "\n(recommended)" if x == pick else "")
                        for x in W.OFFERS], fontsize=10.5, color="#1d1d1b")
    ax.set_yticks(range(len(rows)))
    ax.set_yticklabels(rows, fontsize=10.5, color="#1d1d1b")
    ax.get_yticklabels()[0].set_fontweight("bold")
    ax.xaxis.tick_top()
    ax.tick_params(length=0)
    for sp in ax.spines.values():
        sp.set_visible(False)
    ax.axhline(0.5, color="#ffffff", lw=4)
    j = W.OFFERS.index(pick)
    ax.add_patch(Rectangle((j - 0.5, -0.5), 1, len(rows), fill=False, ec="#1d1d1b", lw=2.4))
    figr.suptitle(f"Put the {SHORT[pick]} on every card in 2027: CHF {fig['chain'][pick]:.2f} net margin per "
                  f"active card-week,\nahead of the {SHORT[run]} at CHF {fig['chain'][run]:.2f}",
                  fontsize=13, fontweight="bold", color="#1d1d1b", x=0.06, ha="left", y=0.985)
    figr.text(0.06, 0.02, "Run-rate weeks 13 July to 23 August 2026, net of reward cost, memberships whose cards all "
              "drew the same cell. Regions as from 1 July 2026.", fontsize=8.5, color="#5f5e5a")
    figr.subplots_adjust(left=0.2, right=0.97, top=0.80, bottom=0.07)
    figr.savefig(path, dpi=150, metadata={"Software": None})
    plt.close(figr)


def main(target=None, golden_dir=None, quiet=False):
    target = target or os.path.join(TASK, "target")
    golden_dir = golden_dir or os.path.join(TASK, "golden")
    os.makedirs(golden_dir, exist_ok=True)
    fig = figures(target)
    write_docx(fig, os.path.join(golden_dir, "programme_offer_2027.docx"))
    write_png(fig, os.path.join(golden_dir, "offer_grid_2027.png"))
    import postprocess
    postprocess.normalize_zip(os.path.join(golden_dir, "programme_offer_2027.docx"), when=(2026, 10, 12, 8, 15, 0))
    if not quiet:
        print("pick", fig["pick"], "runner-up", fig["runner_up"], "gap %.4f" % fig["gap"])
        for x in W.OFFERS:
            print(x, "%.4f" % fig["chain"][x], "value %.1f" % fig["value_chf"][x])
        for r, _ in W.REGIONS:
            print(r, {x: round(fig["grid"][r][x], 4) for x in W.OFFERS})
        print("planned total", fig["planned_total"])
    return fig


if __name__ == "__main__":
    main()
