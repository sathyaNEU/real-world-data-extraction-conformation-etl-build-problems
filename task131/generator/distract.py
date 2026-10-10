"""Files in the same world that the decision does not use."""
import datetime as dt
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font

import world as W


def footfall_csv(out_path, sim):
    rng = np.random.default_rng(W.SEED + 23)
    fmt = {"S": 2100, "M": 4300, "L": 7900}
    rows = []
    for b in W.BRANCHES:
        base = fmt[b[5]] * rng.uniform(0.85, 1.15)
        for i, w in enumerate(W.WEEKS):
            closed = b[0] in W.CLOSURES and W.WEEK_START[w] >= dt.date(2026, 7, 1)
            season = [1.0, 1.01, 0.97, 1.03, 0.99, 0.97, 1.02, 1.0, 0.98, 0.97, 0.96, 0.95, 0.92, 0.9,
                      0.88, 0.87, 0.88, 0.9, 0.91, 0.93][i]
            if closed:
                continue
            doors = int(base * season * rng.lognormal(0, 0.04))
            if b[0] == "6101" and W.WEEK_START[w] >= dt.date(2026, 7, 1):
                doors = int(doors * 1.11)
            rows.append({"branch_no": b[0], "iso_week": w, "door_count": doors,
                         "counter_status": "ok" if rng.random() > 0.015 else "partial"})
    pd.DataFrame(rows).to_csv(out_path, index=False)
    return len(rows)


def panel_survey_xlsx(out_path):
    rng = np.random.default_rng(W.SEED + 29)
    wb = openpyxl.Workbook()
    wb.properties.creator = "Marktforschung Ostschweiz AG"
    wb.properties.created = dt.datetime(2026, 7, 6, 11, 0)
    wb.properties.title = "Talwerk Mitgliederpanel Juni 2026"
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Talwerk member panel, June 2026 wave (fieldwork 8 to 21 June, n = 1,214 members)"
    ws["A1"].font = Font(bold=True)
    hdr = ["Region", "n", "Satisfied with the member programme (%)", "Uses the app weekly (%)",
           "Noticed a new offer in June (%)", "Shops mainly at one Talwerk branch (%)"]
    for j, h in enumerate(hdr, 1):
        ws.cell(row=3, column=j, value=h).font = Font(bold=True)
    ns = [262, 231, 198, 189, 148, 186]
    for i, ((code, name), n) in enumerate(zip(W.REGIONS, ns), 4):
        vals = [name, n, round(rng.uniform(71, 84), 1), round(rng.uniform(38, 52), 1),
                round(rng.uniform(44, 63), 1), round(rng.uniform(66, 81), 1)]
        for j, v in enumerate(vals, 1):
            ws.cell(row=i, column=j, value=v)
    ws.cell(row=11, column=1, value="Regions as in the panel sample frame (January 2026 branch register).")
    ws.cell(row=12, column=1, value="Weighted to the membership by region and age band. Margin of error +/- 2.8 points overall.")
    for col, wdt in zip("ABCDEF", (22, 8, 36, 24, 30, 36)):
        ws.column_dimensions[col].width = wdt
    a = wb.create_sheet("By age band")
    bands = ["18-29", "30-44", "45-59", "60-74", "75+"]
    a.append(["Age band", "n", "Satisfied (%)", "App weekly (%)", "Noticed a new offer (%)"])
    for bnd, n in zip(bands, [141, 302, 351, 298, 122]):
        a.append([bnd, n, round(rng.uniform(68, 86), 1), round(rng.uniform(12, 71), 1), round(rng.uniform(40, 66), 1)])
    wb.save(out_path)


def competitor_checks_csv(out_path):
    rng = np.random.default_rng(W.SEED + 31)
    comps = ["Migros", "Coop", "Denner", "Aldi Suisse", "Lidl Schweiz"]
    rows = []
    for b in W.BRANCHES:
        near = rng.choice(comps, size=rng.integers(1, 4), replace=False)
        for month in ["2026-04", "2026-05", "2026-06", "2026-07", "2026-08"]:
            for c in near:
                idx = {"Migros": 99.0, "Coop": 101.5, "Denner": 92.0, "Aldi Suisse": 86.5, "Lidl Schweiz": 87.0}[c]
                rows.append({"branch_no": b[0], "month": month, "competitor": c,
                             "basket_index_vs_talwerk": round(idx * rng.lognormal(0, 0.012), 1),
                             "items_checked": int(rng.integers(38, 61))})
    pd.DataFrame(rows).to_csv(out_path, index=False)
    return len(rows)
