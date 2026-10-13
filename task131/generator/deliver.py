"""task131 generator: the golden deliverables and generator/figures.json.

tariff_default_advice_2027-28.pdf carries the call, each design's change to the 2027-28 50% POE maximum demand
and a chart of the three changes; tariff_trial_tables_2025-26.xlsx carries the three asks. Every figure comes from
the record the checks wrote, never typed."""
import io
import json
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt   # noqa: E402
from openpyxl import Workbook      # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill   # noqa: E402
from openpyxl.utils import get_column_letter               # noqa: E402
from reportlab.lib import colors   # noqa: E402
from reportlab.lib.pagesizes import A4   # noqa: E402
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet   # noqa: E402
from reportlab.lib.units import mm   # noqa: E402
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle   # noqa: E402

import numpy as np          # noqa: E402
import params as P          # noqa: E402
import postprocess          # noqa: E402

NAMES = {"A": "Design A (Evening 3 to 9)", "B": "Design B (Evening 4 to 8)", "C": "Design C (Stepped)"}
HALF = {41: "8.00 to 8.30 pm", 43: "9.00 to 9.30 pm", 37: "6.00 to 6.30 pm"}


def mw(x):
    return int(round(x))


def lowers(x):
    v = mw(x)
    return "lowers it by %d MW" % v if v > 0 else ("raises it by %d MW" % -v if v < 0 else "leaves it unchanged")


def figures(rec):
    ch = rec["change"]
    ti = rec["ti"]
    f = {
        "call": "C",
        "call_sentence": "Move every flat-rate residential customer onto Design C (Stepped) from 1 July 2027: it lowers the 2027-28 summer maximum demand at 50%% probability of exceedance by %d MW." % mw(ch["C"]),
        "change_mw": {a: mw(ch[a]) for a in P.DESIGNS},
        "md50_mw": round(rec["md50"], 1),
        "n_flat": rec["n_flat"],
        "hot_workdays": rec["hot_days"],
        "hot_tmax_range": [rec["hot_lo"], rec["hot_hi"]],
        "md_half_hour": {a: HALF[ti[a]] for a in P.DESIGNS},
        "default_B_mw": mw(rec["R0"]["B"]),
        "rung2_A_mw": mw(rec["R2"]["A"]),
        "poe": 50,
        "n_hot": len(rec["hot_days"]),
        "design_tmax": round(rec["design_tmax"], 1),
        "mild_max": round(rec["mild_max"], 1),
        "pulldown_kw": {a: round(rec["pull"][a], 1) for a in ("A", "B")},
        "retention": rec["ret"],
        "bills": {a: {"more": rec["bills"][a]["more"], "mean": round(rec["bills"][a]["mean"], 2),
                      "mean_as_written": "%.2f" % abs(rec["bills"][a]["mean"])} for a in P.DESIGNS},
        "survey": {a: {"answered": rec["survey"][a]["answered"], "likely": rec["survey"][a]["likely"],
                       "share": round(rec["survey"][a]["share"], 1)} for a in P.DESIGNS},
        # the ladder, fork and closure figures the design note quotes, as checks.py measured them
        "note": _rounded(rec["note"]),
    }
    return f


def _rounded(x):
    if isinstance(x, dict):
        return {str(k): _rounded(v) for k, v in x.items()}
    if isinstance(x, (list, tuple)):
        return [_rounded(v) for v in x]
    if isinstance(x, (float, np.floating)):
        return round(float(x), 3)
    if isinstance(x, (np.integer,)):
        return int(x)
    return x


def chart(f, path_png):
    fig, ax = plt.subplots(figsize=(7.2, 3.1), dpi=200)
    order = ["C", "A", "B"]
    vals = [f["change_mw"][a] for a in order]
    cols = ["#1f6f4a" if a == "C" else "#b4433b" for a in order]
    y = range(len(order))
    ax.barh(list(y), vals, color=cols, height=0.55)
    ax.axvline(0, color="#333333", lw=0.9)
    ax.set_yticks(list(y))
    ax.set_yticklabels([NAMES[a] for a in order], fontsize=8.5)
    ax.invert_yaxis()
    for k, (a, v) in enumerate(zip(order, vals)):
        lab = ("%d MW lower" % v) if v > 0 else ("%d MW higher" % -v)
        ax.text(v + (0.6 if v > 0 else -0.6), k, lab, va="center", ha="left" if v > 0 else "right", fontsize=8.2)
    ax.set_xlim(min(vals) - 8, max(vals) + 9)
    ax.set_xlabel("Change to the 2027-28 summer maximum demand at 50% POE (MW lowered; negative raises it)", fontsize=7.8)
    ax.tick_params(axis="x", labelsize=7.8)
    ax.set_title("Only the stepped design lowers the design-day maximum", fontsize=9.5, loc="left", fontweight="bold")
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    fig.tight_layout()
    fig.savefig(path_png, dpi=200, metadata={"Software": None})
    plt.close(fig)


def advice_pdf(f, path):
    tmp_png = path + ".chart.png"
    chart(f, tmp_png)
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["BodyText"], fontName="Helvetica", fontSize=9.6, leading=13.2, spaceAfter=5)
    head = ParagraphStyle("h", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, spaceBefore=7, spaceAfter=3)
    title = ParagraphStyle("t", parent=ss["Title"], fontName="Helvetica-Bold", fontSize=13.5, alignment=0, leading=17, spaceAfter=6)
    small = ParagraphStyle("s", parent=body, fontSize=7.8, leading=10, textColor=colors.HexColor("#444444"))
    ch = f["change_mw"]
    lo, hi = f["hot_tmax_range"]
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=17 * mm,
                            bottomMargin=17 * mm, title="2027-28 residential default tariff", author="Pricing and Tariffs",
                            creator="Pricing and Tariffs", invariant=1)
    cell = ParagraphStyle("c", parent=body, fontSize=8.6, leading=10.5, spaceAfter=0)
    cellb = ParagraphStyle("cb", parent=cell, fontName="Helvetica-Bold")
    rows = [[Paragraph(x, cellb) for x in ("Design", "Change to the 2027-28 maximum demand at 50% POE", "When the design-day maximum falls")]]
    for a in ("C", "A", "B"):
        rows.append([Paragraph(x, cell) for x in (NAMES[a], ("%d MW lower" % ch[a]) if ch[a] > 0 else ("%d MW higher" % -ch[a]), f["md_half_hour"][a])])
    tb = Table(rows, colWidths=[52 * mm, 64 * mm, 54 * mm])
    tb.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#9a9a9a")),
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6eaef")),
                            ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#e7f2ec"))]))
    flow = [
        Paragraph("Murrandi Regional Power, Pricing and Tariffs. Advice for the 2027-28 residential tariff schedule, "
                  "9 October 2026. Prepared by Pamela Crawford.", small),
        Spacer(1, 4),
        Paragraph("Default Home Plan households onto Design C (Stepped) from 1 July 2027: it lowers the 2027-28 design-day "
                  "maximum by %d MW, and Designs A and B would raise it" % ch["C"], title),
        Paragraph("Moving every flat-rate (Home Plan) residential customer onto Design C from 1 July 2027, with the same "
                  "right to return to the Home Plan that trial households had, lowers the 2027-28 summer maximum demand at "
                  "50%% probability of exceedance by %d MW. Design A would raise it by %d MW and Design B by %d MW. "
                  "Keeping the Home Plan leaves it at the forecast's %.1f MW." % (ch["C"], -ch["A"], -ch["B"], f["md50_mw"]), body),
        tb,
        Spacer(1, 6),
        Image(tmp_png, width=170 * mm, height=73 * mm),
        Paragraph("Why the two evening designs raise the maximum", head),
        Paragraph("Households on a design set their air conditioners up through the higher-priced hours and pull the "
                  "heat back out when the price falls. On the hottest trial workdays that pull-down, every responding "
                  "home at once, was larger than the cut it repaid. Added to the forecast's design day, it moves the "
                  "maximum out of the window: to the half-hour from 8 pm under Design B and from 9 pm under Design A, both "
                  "above today's maximum. Design C's shoulders split the return between 8 pm and 10 pm, and its maximum "
                  "falls at 8 pm, below the forecast's.", body),
        Paragraph("How the changes were worked out", head),
        Paragraph("Each design's effect is its households' demand less the control households', as assigned (households "
                  "that returned to the Home Plan or left stay in their design's count), for every half-hour of the trial "
                  "workdays at the design day's heat: %d workdays at %.1f to %.1f degrees, against the design day's %.1f, "
                  "with no other workday above %.1f. Each half-hour's effect is multiplied by the forecast's %s Home Plan "
                  "customers for 2027-28 and added to the 2027-28 50%% POE design-day profile; the change is the forecast "
                  "maximum less the highest half-hour of the new profile."
                  % (len(f["hot_workdays"]), lo, hi, f["design_tmax"], f["mild_max"], "{:,}".format(f["n_flat"])), body),
        Paragraph("Changes are to the nearest megawatt. The trial's ordinary workdays do not stand for the design day: on "
                  "them the pull-down after the window is small, and reading the effect there, or only inside the "
                  "windows, would rank Design A or Design B first.", small),
    ]

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(20 * mm, 9 * mm, "Murrandi Regional Power, Pricing and Tariffs: 2027-28 residential default tariff")
        canvas.drawRightString(190 * mm, 9 * mm, "Page %d" % d.page)
        canvas.restoreState()
    doc.build(flow, onFirstPage=footer, onLaterPages=footer)
    os.remove(tmp_png)
    postprocess.scrub(path, P.ORG, "2026-10-09")


def tables_xlsx(f, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Households on 28 Feb 2026"
    hdr = ["Design", "Still billed on the design", "Back on the Home Plan at own request", "No longer our customers", "Assigned"]
    _hdr(ws, hdr)
    for r, a in enumerate(P.DESIGNS, start=2):
        x = f["retention"][a]
        for j, v in enumerate([NAMES[a], x["still"], x["back"], x["gone"], P.PER_ARM], start=1):
            ws.cell(row=r, column=j, value=v)
    _w(ws, [28, 16, 20, 16, 10])
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Charges Dec 2025 to Feb 2026")
    _hdr(ws, ["Design", "Households charged more than the Home Plan", "Mean difference per household ($)", "Households"])
    for r, a in enumerate(P.DESIGNS, start=2):
        x = f["bills"][a]
        for j, v in enumerate([NAMES[a], x["more"], x["mean"], P.PER_ARM], start=1):
            c = ws.cell(row=r, column=j, value=v)
            if j == 3:
                c.number_format = "#,##0.00;-#,##0.00"
    _w(ws, [28, 22, 22, 12])
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("March 2026 survey")
    _hdr(ws, ["Design", "Households answering Q4", "Likely or very likely to stay", "Share (%)"])
    for r, a in enumerate(P.DESIGNS, start=2):
        x = f["survey"][a]
        for j, v in enumerate([NAMES[a], x["answered"], x["likely"], x["share"]], start=1):
            c = ws.cell(row=r, column=j, value=v)
            if j == 4:
                c.number_format = "0.0"
    _w(ws, [28, 16, 18, 10])
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Notes")
    notes = [
        "Summer pricing trial 2025-26: tables for the 2027-28 tariff advice",
        "Households are the trial register's households, followed by customer number across their accounts.",
        "Households on 28 Feb 2026: the tariff on each household's general supply that day in the billing system's account history as at 12 March 2026. "
        "Households moved back after a life-support registration are in none of the three columns.",
        "Charges: general supply, 1 December 2025 to 28 February 2026, at the gazetted 2025-26 charges including GST, on the days each household was billed on its design, "
        "against the Home Plan for the same reads and days. Controlled-load hot water and export credits are on the same tariff either way. "
        "A negative difference is a saving.",
        "Survey: each household's first submitted response; telephone answers to Q4 are coded in the order read (very unlikely first).",
    ]
    for r, t in enumerate(notes, start=1):
        c = ws.cell(row=r, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if r == 1:
            c.font = Font(bold=True)
    _w(ws, [130])
    wb.save(path)
    postprocess.scrub(path, "Pamela Crawford", "2026-10-09")
    postprocess.finish_office(path, "2026-10-08T01:12:00Z", "2026-10-09T05:47:00Z", P.ORG)


def _hdr(ws, cols):
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE4EE")
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _w(ws, widths):
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w


def write(rec, golden_dir, figures_path):
    os.makedirs(golden_dir, exist_ok=True)
    f = figures(rec)
    advice_pdf(f, os.path.join(golden_dir, "tariff_default_advice_2027-28.pdf"))
    tables_xlsx(f, os.path.join(golden_dir, "tariff_trial_tables_2025-26.xlsx"))
    with open(figures_path, "w") as fh:
        json.dump(f, fh, indent=1, sort_keys=True)
    return golden_dir, f
