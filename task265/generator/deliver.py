"""The goldens for task265: hip_construct_scorecard_2027.xlsx (each construct's ten-year cumulative percent revision
and whether it stays, the 2026 primary total hips by construct and hospital, the 2026 hips missing from the registry)
and construct_review_memo_2027.docx (the one-page memo to the Value Analysis Committee with the revision curves), plus
generator/figures.json and the task's metadata.json. Every figure comes from analysis.py run on the shipped pack."""
import json
import os
import shutil
import tempfile

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt                          # noqa: E402
import numpy as np                                       # noqa: E402
from openpyxl import Workbook                            # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

import analysis as A                                     # noqa: E402
import common as K                                       # noqa: E402
import postprocess as PP                                 # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
DELIVERABLES = ["hip_construct_scorecard_2027.xlsx", "construct_review_memo_2027.docx"]
CHART = "cpr_curves.png"
COLOURS = {"TA": "#b8442f", "TS": "#d9822b", "MF": "#2f5d50", "VS": "#3d6fa3", "MS": "#6b8f3a", "BP": "#7a3f7d"}


def r1(x):
    return float("%.1f" % (np.floor(x * 10 + 0.5) / 10))


def results(target):
    fig, hips = A.call(target)
    stays = [c for c in K.CODES if fig[c]["golden"] <= K.BAR]
    asks = A.asks(target)
    natural = A.asks(target, honour_side=False, merges=False)
    curves = {}
    for c in K.CODES:
        rows = np.array(A.rows_for(hips, K.CNAME[c], "golden"), float)
        e, s, d = rows[:, 0], rows[:, 1], rows[:, 2].astype(bool)
        ts, S, xs, ys = np.unique(s[d & (s <= K.TEN)]), 1.0, [0.0], [0.0]
        for t in ts:
            at = (e < t) & (s >= t)
            S *= 1 - (d & (s == t) & at).sum() / at.sum()
            xs.append(t / K.YEAR)
            ys.append(100 * (1 - S))
        xs.append(10.0)
        ys.append(ys[-1])
        curves[c] = (xs, ys)
    return dict(fig=fig, stays=stays, asks=asks, natural=natural, curves=curves)


# ---------------------------------------------------------------------------------------------- the scorecard
def _head(ws, head, widths):
    ws.append(head)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w


def workbook(path, o):
    F, a = o["fig"], o["asks"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Scorecard"
    n_ok = len(o["stays"])
    ws.append(["%s of the six hip constructs meet the ten-year revision bar for the 1 July 2027 hip agreement" %
               {3: "Three"}.get(n_ok, str(n_ok))])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Policy VA-07, section 4. Brennick Valley Joint Registry extract of 31 December 2026; 2026 hips from "
               "the perioperative implant log. Definitions on the Notes sheet."])
    ws["A2"].font = Font(italic=True, size=9, color="5B646B")
    ws.append([])
    _head(ws, ["Construct", "Revised by 10 years (%)", "Stays on the agreement", "2026 hips, Brennick Valley Medical Center",
               "2026 hips, Tollefson Lake Hospital", "2026 hips, Kittelstad Orthopaedic Hospital", "2026 hips, all three",
               "2026 hips missing from the registry"], [19, 12, 12, 15, 15, 15, 11, 14])
    ws.row_dimensions[4].height = 44
    for c in K.CODES:
        row = [K.CNAME[c], r1(F[c]["golden"]), "Yes" if c in o["stays"] else "No"] + [a["A"][c][h] for h in K.HOSP] + \
            [sum(a["A"][c].values()), a["C_construct"][c]]
        ws.append(row)
        r = ws.max_row
        ws.cell(r, 2).number_format = "0.0"
        for col in range(2, 9):
            ws.cell(r, col).alignment = Alignment(horizontal="right" if col != 3 else "center")
        for col in range(4, 9):
            ws.cell(r, col).number_format = "#,##0"
        if c not in o["stays"]:
            ws.cell(r, 3).font = Font(bold=True, color="9C2B1E")
    ws.append(["All six", None, None] + [sum(a["A"][c][h] for c in K.CODES) for h in K.HOSP] +
              [sum(sum(a["A"][c].values()) for c in K.CODES), sum(a["C_construct"].values())])
    r = ws.max_row
    for col in range(1, 9):
        ws.cell(r, col).font = Font(bold=True)
        if col >= 4:
            ws.cell(r, col).number_format = "#,##0"
            ws.cell(r, col).alignment = Alignment(horizontal="right")
    ws.freeze_panes = "B5"
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "landscape"

    ws2 = wb.create_sheet("Registry capture 2026")
    ws2.append(["2026 primary total hips missing from the registry, by hospital"])
    ws2["A1"].font = Font(bold=True, size=12)
    ws2.append([])
    _head(ws2, ["Hospital", "2026 primary total hips", "Missing from the registry"], [34, 14, 14])
    for h in K.HOSP:
        ws2.append([K.HOSP_NAME[h], sum(a["A"][c][h] for c in K.CODES), a["C_hosp"][h]])
        for col in (2, 3):
            ws2.cell(ws2.max_row, col).number_format = "#,##0"
    ws2.append(["All three", sum(a["A"][c][h] for c in K.CODES for h in K.HOSP), sum(a["C_hosp"].values())])
    for col in range(1, 4):
        ws2.cell(ws2.max_row, col).font = Font(bold=True)
    ws2.freeze_panes = "A4"

    ws3 = wb.create_sheet("Notes")
    ws3.column_dimensions["A"].width = 118
    for line in [
        "Construct: one femoral stem family with one acetabular cup family (Policy VA-07, 4.1).",
        "Revised by 10 years: cumulative percent revision 3,652.5 days after the operation, Kaplan-Meier on the "
        "registry's hips with the construct, first revision of any component as the event; deaths (Wisconsin death "
        "index linkage) and moves out of Wisconsin censored (4.3). Stays: 5.0 per cent or less (4.2).",
        "Hips replaced from 2005 to 2017 were entered at their patient's first registry clinic visit (2018 to 2021); "
        "each is counted at risk from that visit. Hips captured from 2018 are at risk from the operation.",
        "Death dates the linkage gives as month and year only are taken as the 15th; the 1st or the last day of the "
        "month changes no figure at one decimal.",
        "2026 hips: primary total hip cases in the 2026 perioperative implant log, one hip for each side the case "
        "list records as operated on, with an implanted stem and shell; components recorded for the other side of a "
        "one-sided case are the patient's existing hip and are not counted.",
        "Missing from the registry: no 2026 row in the registry's prospective export with the same medical record "
        "number (or a number HIM merged into it in 2026), side and operation date.",
    ]:
        ws3.append([line])
        ws3.cell(ws3.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


# ---------------------------------------------------------------------------------------------- the chart and memo
def chart(path, o):
    fig, ax = plt.subplots(figsize=(7.2, 3.55), dpi=150)
    order = sorted(K.CODES, key=lambda c: -o["fig"][c]["golden"])
    ends = sorted(K.CODES, key=lambda c: o["curves"][c][1][-1])
    ly, last = {}, -9.0
    for c in ends:
        ly[c] = max(o["curves"][c][1][-1], last + 0.36)
        last = ly[c]
    for c in order:
        xs, ys = o["curves"][c]
        stay = c in o["stays"]
        ax.step(xs, ys, where="post", color=COLOURS[c], lw=1.2 if stay else 1.9, ls="-",
                alpha=1.0 if not stay else .9, label=K.CNAME[c])
        ax.text(10.08, ly[c], "%s %.1f%%" % (K.CNAME[c], r1(o["fig"][c]["golden"])), va="center", fontsize=6.8,
                color=COLOURS[c], fontweight="bold" if not stay else "normal")
    ax.axhline(K.BAR, color="#1d2326", lw=.9, ls=":")
    ax.text(.12, K.BAR + .14, "Policy VA-07 bar, 5.0% revised by 10 years", fontsize=7, color="#1d2326")
    ax.set_xlim(0, 10)
    ax.set_ylim(0, max(max(v[1]) for v in o["curves"].values()) + .8)
    ax.set_xticks(range(0, 11))
    ax.set_xlabel("Years after the operation", fontsize=8)
    ax.set_ylabel("Hips revised (cumulative %)", fontsize=8)
    ax.tick_params(labelsize=7.5)
    off = [K.CNAME[c] for c in K.CODES if c not in o["stays"]]
    ax.set_title("%s, %s and %s cross the bar before ten years" % tuple(off), fontsize=8.6, loc="left")
    ax.grid(color="#e7ebe6", lw=.6)
    ax.set_axisbelow(True)
    for s_ in ("top", "right"):
        ax.spines[s_].set_visible(False)
    fig.subplots_adjust(left=.075, right=.80, top=.9, bottom=.14)
    fig.savefig(path, metadata={"Software": None})
    plt.close(fig)


def memo(path, o, png):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt, RGBColor
    F, a = o["fig"], o["asks"]
    GREY = RGBColor(0x5B, 0x64, 0x6B)
    doc = docx.Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(9.5)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(215.9), Mm(279.4)
    sec.left_margin = sec.right_margin = Mm(19)
    sec.top_margin, sec.bottom_margin = Mm(13), Mm(12)
    hp = sec.header.paragraphs[0]
    hp.text = "Brennick Valley Health | Supply Chain, Value Analysis | Committee meeting of 13 May 2027, item 4"
    hp.runs[0].font.size = Pt(7.5)
    hp.runs[0].font.color.rgb = GREY

    def para(text="", size=None, bold=False, color=None, after=3, before=0):
        p = doc.add_paragraph()
        if text:
            r = p.add_run(text)
            r.bold = bold
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        return p

    for k, v in (("To", "Value Analysis Committee (chair: %s, Director of Value Analysis)" % K.PEOPLE["value"]),
                 ("From", "%s, orthopaedic outcomes analyst" % K.PEOPLE["analyst"]),
                 ("Date", "23 April 2027"),
                 ("Re", "Hip agreement renewal, 1 July 2027: constructs under Policy VA-07, section 4")):
        p = para(after=0)
        r = p.add_run("%s:\t" % k)
        r.bold = True
        p.add_run(v)
    stays = [K.CNAME[c] for c in o["stays"]]
    off = [K.CNAME[c] for c in K.CODES if c not in o["stays"]]
    para("Three of the six hip constructs meet the ten-year revision bar", 13, bold=True, before=8, after=4)
    p = para(after=4)
    p.add_run("Keep %s, %s and %s on the agreement. Take %s, %s and %s off it for primary operations; under "
              "section 4.5 their components stay available for revisions." % (*stays, *off)).bold = True
    tb = doc.add_table(rows=1, cols=4)
    tb.style = "Light List"
    for i, t in enumerate(["Construct", "Revised by 10 years", "Policy VA-07, 4.2", "2026 hips, all three hospitals"]):
        tb.rows[0].cells[i].text = t
    for c in sorted(K.CODES, key=lambda c: F[c]["golden"]):
        cells = tb.add_row().cells
        cells[0].text = K.CNAME[c]
        cells[1].text = "%.1f%%" % r1(F[c]["golden"])
        cells[2].text = "Stays" if c in o["stays"] else "Comes off"
        cells[3].text = "{:,}".format(sum(a["A"][c].values()))
        for i in (1, 3):
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for row in tb.rows:
        for i, w in enumerate((48, 34, 34, 44)):
            row.cells[i].width = Mm(w)
    para("Source: Brennick Valley Joint Registry extract of 31 December 2026, Kaplan-Meier, deaths and moves out of "
         "Wisconsin censored; 2026 hips from the perioperative implant log.", 7.5, color=GREY, before=2, after=6)
    para("Most of the Edvalt constructs' ten-year evidence is hips replaced before the registry began in January "
         "2018. The registry entered those hips at the patient's first registry clinic visit, between 2018 and 2021, "
         "so a hip revised before its visit never reached the registry, and the curves below count each of them at "
         "risk only from the visit. Clocked from the operation instead, the years before the visit enter as "
         "revision-free and both Edvalt constructs come out near 3 per cent (%.1f and %.1f)."
         % (r1(F["TA"]["op"]), r1(F["TS"]["op"])), after=4)
    doc.add_picture(png, width=Mm(172))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tot = {h: sum(a["A"][c][h] for c in K.CODES) for h in K.HOSP}
    para("For the renewal's volumes: the three hospitals put the six constructs into {:,} primary total hips in 2026 "
         "(Brennick Valley Medical Center {}, Tollefson Lake {}, Kittelstad {}); the scorecard has them by construct "
         "and hospital. {} of those hips are not on the registry (Brennick Valley Medical Center {}, Tollefson Lake {}, "
         "Kittelstad {}).".format(
             sum(tot.values()), tot["BVMC"], tot["TLH"], tot["KOH"], sum(a["C_hosp"].values()), a["C_hosp"]["BVMC"],
             a["C_hosp"]["TLH"], a["C_hosp"]["KOH"]), before=4, after=3)
    para("Death dates the state linkage gives as month and year only are taken as the 15th; the 1st or the last day "
         "of the month changes no figure at one decimal. Ten years is 3,652.5 days.", 7.5, color=GREY, after=0)
    doc.save(path)


# ---------------------------------------------------------------------------------------------- figures.json
def figures(o):
    F, a, n = o["fig"], o["asks"], o["natural"]
    return dict(
        call=dict(stays=[K.CNAME[c] for c in o["stays"]], off=[K.CNAME[c] for c in K.CODES if c not in o["stays"]],
                  cpr10={K.CNAME[c]: r1(F[c]["golden"]) for c in K.CODES}, bar=K.BAR),
        rungs={r: {K.CNAME[c]: r1(F[c][r]) for c in K.CODES} for r in ("crude", "op", "reg", "golden")},
        rung_sets={r: [K.CNAME[c] for c in A.keeps(F, r)] for r in ("crude", "op", "reg", "golden")},
        registry=dict(hips={K.CNAME[c]: F[c]["n"] for c in K.CODES}, legacy={K.CNAME[c]: F[c]["legacy"] for c in K.CODES},
                      at_risk_10={K.CNAME[c]: F[c]["at_risk_10"] for c in K.CODES},
                      prospective_max_years={K.CNAME[c]: round(F[c]["prosp_max_years"], 2) for c in K.CODES}),
        ask_a={K.CNAME[c]: {K.HOSP_NAME[h]: a["A"][c][h] for h in K.HOSP} for c in K.CODES},
        ask_a_hospital_totals={K.HOSP_NAME[h]: sum(a["A"][c][h] for c in K.CODES) for h in K.HOSP},
        ask_a_total=sum(a["A"][c][h] for c in K.CODES for h in K.HOSP),
        ask_a_construct_totals={K.CNAME[c]: sum(a["A"][c].values()) for c in K.CODES},
        ask_c_hospital={K.HOSP_NAME[h]: a["C_hosp"][h] for h in K.HOSP},
        ask_c_construct={K.CNAME[c]: a["C_construct"][c] for c in K.CODES},
        ask_c_total=sum(a["C_hosp"].values()),
        natural=dict(ask_a={K.CNAME[c]: {K.HOSP_NAME[h]: n["A"][c][h] for h in K.HOSP} for c in K.CODES},
                     ask_c_hospital={K.HOSP_NAME[h]: n["C_hosp"][h] for h in K.HOSP},
                     ask_c_construct={K.CNAME[c]: n["C_construct"][c] for c in K.CODES}),
        bar=K.BAR, years=10, ten_years_days=3652.5, month_only_death_day=15, policy_sections=["4.2", "4.3"])


def write(target, golden, T=None):
    os.makedirs(golden, exist_ok=True)
    o = results(target)
    xp = os.path.join(golden, DELIVERABLES[0])
    workbook(xp, o)
    tmp = tempfile.mkdtemp()
    png = os.path.join(tmp, CHART)
    chart(png, o)
    dp = os.path.join(golden, DELIVERABLES[1])
    memo(dp, o, png)
    shutil.rmtree(tmp)
    PP.COMPANY = K.SYSTEM
    PP.normalize_ooxml(xp, K.PEOPLE["analyst"], K.PEOPLE["analyst"], "2027-04-19T14:02:37Z", "2027-04-23T20:41:15Z",
                       excel=True)
    PP.normalize_ooxml(dp, K.PEOPLE["analyst"], K.PEOPLE["analyst"], "2027-04-20T15:11:05Z", "2027-04-23T21:02:44Z",
                       title="Hip constructs for the 1 July 2027 agreement")
    Fj = figures(o)
    json.dump(Fj, open(os.path.join(HERE, "figures.json"), "w"), indent=1, sort_keys=True)
    meta = {"distractor_files": ["target/" + K.FILES[k] for k in K.DISTRACTORS], "stated_inputs": ["5.0"]}
    json.dump(meta, open(os.path.join(os.path.dirname(target), "metadata.json"), "w"), indent=2)
    if T is not None:
        T["o"] = o
    return o, Fj
